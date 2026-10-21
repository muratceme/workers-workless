import json
import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
import tablo as T  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def play(tur, tutar, gun=1, sku="plus_aylik", para="TRY"):
    return {"platform": "Google Play", "tarih": date(2026, 9, gun), "tur": tur, "ham_tur": tur, "siparis": "S1", "urun": sku, "urun_adi": sku, "ulke": "TR",
            "para": para, "tutar": D(tutar), "adet": 1 if tur in ("Satış", "İade") else 0, "abonelik": ""}


def apple(adet, pay, para="TRY", abonelik="New", gun=1):
    return {"platform": "App Store", "tarih": date(2026, 9, gun), "tur": "İade" if adet < 0 else "Satış", "ham_tur": "IAY", "siparis": "", "urun": "plus.aylik",
            "urun_adi": "Plus", "ulke": "US", "para": para, "tutar": D(adet) * D(pay), "adet": abs(adet), "abonelik": abonelik}


class AnalizTesti(unittest.TestCase):
    def test_play_turleri_ve_net(self):
        r = main.analiz([play("Satış", "100"), play("Komisyon", "-15"), play("İade", "-100"), play("Komisyon iadesi", "15"), play("Diğer", "5")], "TRY", {})
        k = r["gunluk"][(date(2026, 9, 1), "Google Play")]
        self.assertEqual((k["brut"], k["iade"], k["komisyon"], k["diger"], k["net"], k["satis_adet"], k["iade_adet"]),
                         (D(100), D(-100), D(0), D(5), D(5), 1, 1))
        self.assertEqual(r["toplam"]["komisyon_orani"], D(0))

    def test_apple_iade_ve_kur(self):
        r = main.analiz([apple(2, "70", "USD"), apple(-1, "70", "USD", ""), apple(1, "3", "EUR")], "TRY", {"USD": D("40")})
        k = r["gunluk"][(date(2026, 9, 1), "App Store")]
        self.assertEqual((k["net"], k["iade"], k["satis_adet"], k["iade_adet"], k["yeni"], k["brut"], k["komisyon"]),
                         (D(2800), D(-2800), 2, 1, 2, D(0), D(0)))       # App Store'da brüt / komisyon hesaplanmaz
        self.assertEqual([(u["onem"], u["kim"]) for u in r["uyarilar"]], [("Yüksek", "EUR")])   # kuru yok → toplama girmez

    def test_iptal_olaylari(self):
        olay = [{"tarih": date(2026, 9, 1), "olay": "İptal", "adet": 2, "platform": "Google Play"}, {"tarih": date(2026, 9, 1), "olay": "Yeni", "adet": 1, "platform": "Google Play"}]
        r = main.analiz([play("Satış", "10")], "TRY", {}, olay)
        self.assertEqual(r["toplam"]["iptal"], 2)
        self.assertEqual([main._olay_turu(x) for x in ("İptal", "Cancel", "Renewal", "Yeni abonelik", "Refund", "Bilinmeyen")],
                         ["İptal", "İptal", "Yenileme", "Yeni", "İade", "Diğer"])


class OkumaTesti(unittest.TestCase):
    def test_play_utf16_ve_bilinmeyen_tur(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "earnings.csv"
            y.write_bytes(("Description,Transaction Date,Transaction Type,Product Title,SKU ID,Buyer Country,Merchant Currency,Amount (Merchant Currency)\n"
                           "GPA.1,\"Sep 2, 2026\",Charge,Plus,plus_aylik,TR,TRY,149.99\nGPA.1,\"Sep 2, 2026\",Tax,Plus,plus_aylik,TR,TRY,-1.00\n").encode("utf-16"))
            u = []
            i = main.oku_play([y], u)
            self.assertEqual([(x["tarih"], x["tur"], x["tutar"]) for x in i], [(date(2026, 9, 2), "Satış", D("149.99")), (date(2026, 9, 2), "Diğer", D("-1.00"))])
            self.assertEqual([x["tur"] for x in u], ["Tanınmayan Play işlem türü"])

    def test_apple_abd_tarihi_ve_ucretsiz_satir(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "s.txt"
            y.write_text("SKU\tTitle\tUnits\tDeveloper Proceeds\tBegin Date\tCountry Code\tCurrency of Proceeds\tSubscription\n"
                         "app\tUygulama\t40\t0\t09/03/2026\tTR\tTRY\t\nplus\tPlus\t-1\t106.25\t09/03/2026\tTR\tTRY\t\n", encoding="utf-8")
            i = main.oku_apple([y], [])
            self.assertEqual([(x["tarih"], x["tur"], x["tutar"]) for x in i], [(date(2026, 9, 3), "İade", D("-106.25"))])   # ücretsiz indirme atlandı


class UcuncuTesti(unittest.TestCase):
    def test_ornek_veri(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "r.xlsx"
            r = main.calistir([ORNEK / "play_earnings_202609.csv"], [ORNEK / "appstore_satis_202609.txt"], cikti, None, ORNEK / "kur.csv", ORNEK / "abonelik_olaylari.csv")
            t = r["toplam"]
            self.assertEqual((r["hedef"], t["net"], t["iptal"], t["satis_adet"], t["iade_adet"]), ("TRY", D("37657.92"), 37, 189, 5))
            self.assertEqual({p: v for p, v in r["platform_net"].items()}, {"Google Play": D("22864.34"), "App Store": D("14793.58")})
            self.assertEqual(T.yuvarla(t["komisyon_orani"] * 100, "0.1"), D("14.8"))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Günlük", "Aylık", "Ürünler", "Ülkeler", "Abonelik Olayları", "İşlemler", "Uyarılar"])
            o = json.loads((Path(tmp) / f"{main.PAKET}.ozet.json").read_text(encoding="utf-8"))
            self.assertEqual((o["paket"], o["para_birimi"], len(o["gunluk"])), (main.PAKET, "TRY", 30))
            self.assertAlmostEqual(sum(g["net"] for g in o["gunluk"].values()), 37657.92, places=2)

    def test_komut_satiri(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "r.xlsx")]), 0)
            self.assertEqual(main.main(["--play", str(Path(tmp) / "yok.csv"), "--cikti", str(Path(tmp) / "r.xlsx")]), 1)


if __name__ == "__main__":
    unittest.main()
