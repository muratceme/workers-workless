import json
import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
G0 = date(2026, 9, 1)


def s(gun, kullanici, olay, adet=1):
    return {"gun": G0 + timedelta(days=gun), "kullanici": kullanici, "olay": olay, "ilk": None, "platform": "ANDROID", "surum": "1.0", "adet": adet}


def veri():
    return [s(0, "A", "first_open"), s(0, "A", "session_start"), s(1, "A", "session_start"), s(7, "A", "tarot", 2),
            s(0, "B", "first_open"), s(3, "B", "session_start"),
            s(1, "C", "first_open"), s(1, "C", "tarot"), s(2, "C", "session_start"),
            s(1, "P", "notification_receive"),          # yalnız bildirim aldı → aktif değil
            s(5, "E", "session_start")]                 # first_open yok → kohort ilk görülme günü


class AnalizTesti(unittest.TestCase):
    def test_dau_kohort_tutma(self):
        r = main.analiz(veri())
        self.assertEqual(r["bitis"], G0 + timedelta(days=7))
        self.assertEqual([r["gunluk"][G0 + timedelta(days=i)]["dau"] for i in range(8)], [2, 2, 1, 1, 0, 1, 0, 1])
        self.assertEqual([r["gunluk"][G0 + timedelta(days=i)]["yeni"] for i in range(8)], [2, 1, 0, 0, 0, 1, 0, 0])
        t0, t1 = r["tutma"][G0], r["tutma"][G0 + timedelta(days=1)]
        self.assertEqual((t0["kullanici"], t0[1], t0[3], t0[7], t0[14]), (2, D("0.5"), D("0.5"), D("0.5"), None))
        self.assertEqual((t1[1], t1[7]), (D(1), None))                                  # D7 günü dökümün dışında
        self.assertEqual(r["ortalama"][1]["oran"], D("0.5"))     # kohortlar 1.9 (2 kişi, 0,5), 2.9 (1 kişi, 1), 6.9 (E, 0): (1 + 1 + 0) ÷ 4
        self.assertEqual((r["ortalama"][7]["oran"], r["ortalama"][7]["kohort"]), (D("0.5"), 1))
        self.assertEqual((r["wau"], r["mau"], r["dau"]), (4, 4, 1))
        self.assertEqual(r["kohort_kaynagi"], {"first_open": 3, "ilk görülme": 1})
        self.assertEqual([u["tur"] for u in r["uyarilar"] if u["onem"] == "Orta"][0], "Kohort ilk görülme gününe göre")

    def test_ozellikler_otomatik_olaylari_saymaz(self):
        r = main.analiz(veri())
        self.assertEqual([(x["ad"], x["adet"], x["kullanici"], x["son7"], x["onceki7"]) for x in r["ozellikler"]], [("tarot", 3, 2, 3, 0)])

    def test_aktif_olmayan_gun_uyarisi(self):
        r = main.analiz(veri())
        self.assertEqual([u["kim"] for u in r["uyarilar"] if u["tur"] == "Veri olmayan gün"], ["05.09.2026", "07.09.2026"])


class OrnekVeriTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "k.xlsx"
            r = main.calistir([ORNEK / "firebase_olay_ozeti.csv"], cikti)
            self.assertEqual((r["ilk"], r["bitis"]), (date(2026, 8, 17), date(2026, 9, 30)))
            gunluk_son = r["gunluk"][date(2026, 9, 30)]
            self.assertEqual(gunluk_son["dau"], r["dau"])
            self.assertEqual(r["kohort_kaynagi"], {"first_open": 220})                  # pasif kullanıcı kohorta girmez
            for n in (1, 7, 30):
                self.assertTrue(D(0) < r["ortalama"][n]["oran"] < D(1))
            self.assertGreater(r["ortalama"][1]["oran"], r["ortalama"][7]["oran"])
            self.assertGreater(r["ortalama"][7]["oran"], r["ortalama"][30]["oran"])
            self.assertEqual([x["ad"] for x in r["ozellikler"]][:3], ["kahve_fali", "burc_yorumu", "tarot"])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Günlük", "Tutma", "Özellikler", "Platform ve Sürüm", "Uyarılar"])
            o = json.loads((Path(tmp) / f"{main.PAKET}.ozet.json").read_text(encoding="utf-8"))
            self.assertEqual((len(o["gunluk"]), sorted(o["tutma"])), (45, ["d1", "d30", "d7"]))

    def test_komut_satiri(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "k.xlsx")]), 0)
            self.assertEqual(main.main(["--tarih", "31/31/2026"]), 2)


if __name__ == "__main__":
    unittest.main()
