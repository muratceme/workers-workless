"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import json
import mailbox
import os
import re
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    cevap = []
    for i in re.findall(r'<eposta id="([^"]+)"', kullanici):
        # Model bilerek hata yapıyor: iadeyi "Abonelik / ödeme" sayıyor ve iade sözü veriyor; kod kuralları düzeltmeli
        k, a, taslak = "Kullanım sorusu", "dusuk", "Merhaba, ekibimiz dönüş yapacak."
        if i == "D-1002":
            k, taslak = "Abonelik / ödeme", "Merhaba, ücretiniz iade edilecek."
        if i == "D-1004":
            k, a = "Hata / çökme", "orta"
        if i == "D-1007":
            k, a = "Giriş / hesap", "yuksek"
        cevap.append({"id": i, "kategori": k, "duygu": "notr", "aciliyet": a, "ozet": "o", "cevap_taslagi": taslak})
    return {"epostalar": cevap}


class OkumaTesti(unittest.TestCase):
    def test_csv_ve_kurallar(self):
        e = {x["id"]: x for x in agent.epostalari_oku([ORNEK / "destek_epostalari.csv"])}
        self.assertEqual(len(e), 12)
        self.assertNotIn("wrote:", e["D-1004"]["mesaj"])                     # alıntılanan eski yazışma atıldı
        self.assertTrue(e["D-1004"]["mesaj"].endswith("yine aynı."))
        bek = {"D-1001": ["odeme_acilmadi"], "D-1002": ["iade"], "D-1003": ["silme"], "D-1006": ["yasal"], "D-1009": ["guvenlik"], "D-1010": [], "D-1005": []}
        for i, k in bek.items():
            self.assertEqual(agent.kurallar(e[i]), k, i)

    def test_eml_ve_mbox(self):
        ham = ("From: Test <t@example.com>\nSubject: =?utf-8?q?Hesab=C4=B1m=C4=B1_silin?=\nDate: Tue, 29 Sep 2026 10:00:00 +0300\n"
               "Content-Type: text/plain; charset=utf-8\n\nHesabımı silin lütfen.\n\n> eski mesaj\n")
        with tempfile.TemporaryDirectory() as tmp:
            (Path(tmp) / "a.eml").write_text(ham, encoding="utf-8")
            e = agent.epostalari_oku([tmp])[0]
            self.assertEqual((e["id"], e["konu"], e["mesaj"], e["zaman"]), ("a", "Hesabımı silin", "Hesabımı silin lütfen.", datetime(2026, 9, 29, 10)))
            mb = mailbox.mbox(Path(tmp) / "k.mbox")
            mb.add(ham.encode("utf-8"))
            mb.add(ham.encode("utf-8"))
            mb.close()
            self.assertEqual([x["id"] for x in agent.epostalari_oku([Path(tmp) / "k.mbox"])], ["M1", "M2"])


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_kod_kurallari(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "r.xlsx"
            s = agent.calistir([ORNEK / "destek_epostalari.csv"], cikti, ORNEK / "sla.json", evet=True)
            r = s["sonuc"]
            gonderilen = "\n".join(MESAJLAR)
            self.assertNotIn("@example.com", gonderilen)                      # gönderen ve adresler modele gitmez
            self.assertNotIn("0532 111 22 33", gonderilen)
            self.assertEqual(r["D-1002"]["kategori"], "İade talebi")          # kural modeli düzeltir
            self.assertTrue(r["D-1002"]["onay"])
            self.assertTrue(r["D-1002"]["soz"])                               # "iade edilecek" yakalandı
            self.assertEqual([r[i]["oncelik"] for i in ("D-1001", "D-1006", "D-1009", "D-1007")], ["P1"] * 4)
            self.assertEqual((r["D-1003"]["oncelik"], r["D-1003"]["kategori"], r["D-1003"]["onay"]), ("P2", "Hesap / veri silme (KVKK)", True))
            self.assertEqual((r["D-1004"]["oncelik"], r["D-1010"]["oncelik"], r["D-1010"]["onay"]), ("P2", "P3", False))
            self.assertEqual(r["D-1001"]["hedef"], datetime(2026, 9, 29, 12, 12))   # P1: 4 saat
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Onay Bekleyenler", "Tüm E-postalar", "Özet"])
            onay = {wb["Onay Bekleyenler"].cell(i, 1).value for i in range(2, wb["Onay Bekleyenler"].max_row + 1)}
            self.assertEqual(onay, {"D-1002", "D-1003", "D-1006", "D-1009"})
            o = json.loads((Path(tmp) / f"{agent.PAKET_ADI}.ozet.json").read_text(encoding="utf-8"))
            self.assertEqual(o["gunluk"]["2026-09-29"], {"talep": 10, "p1": 4, "onay": 4})

    def test_bos_girdi(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "b.csv"
            y.write_text("Tarih,Konu,Mesaj\n", encoding="utf-8")
            with self.assertRaises(llm.LLMHatasi):
                agent.calistir([y], Path(tmp) / "r.xlsx", evet=True)


if __name__ == "__main__":
    unittest.main()
