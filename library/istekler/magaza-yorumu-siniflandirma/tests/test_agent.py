"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import json
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
PLAY = ORNEK / "reviews_reviews_com.ornek.uygulama_202609.csv"
APPLE = ORNEK / "appstore_yorumlari.csv"
MESAJLAR = []
TALEP = {"sesli okuma": "sesli okuma", "karanlık": "karanlık mod", "dark mode": "karanlık mod", "aylık": "aylık ödeme seçeneği"}


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    cevap = []
    for i, metin in re.findall(r'<yorum id="([^"]+)"[^>]*>\n(.*?)\n</yorum>', kullanici, re.S):
        talep = next((v for k, v in TALEP.items() if k in metin.lower()), "")
        alinti = metin.split(".")[0]
        if i == "GP7":
            alinti = "Bu uygulama tamamen bozuk"                  # uydurma alıntı → silinmeli
        cevap.append({"id": i, "kategori": "Diğer", "duygu": "olumlu", "ozellik_talebi": talep, "alinti": alinti,   # hepsine "olumlu": kural düzeltmeli
                      "cevap_taslagi": "x" * 400 if i == "GP5" else "Teşekkürler, ekibimize ilettik."})
    return {"yorumlar": cevap}


class OkumaTesti(unittest.TestCase):
    def test_play_ve_appstore(self):
        y = agent.yorumlari_oku([PLAY], [APPLE])
        self.assertEqual((len(y), sum(x["magaza"] == "Google Play" for x in y)), (24, 18))   # Play UTF-16
        self.assertEqual(sum(1 for x in y if not (x["metin"] or x["baslik"])), 3)            # yalnız puan
        self.assertTrue(next(x for x in y if x["id"] == "GP3")["cevaplandi"])

    def test_play_olmayan_dosya(self):
        with self.assertRaises(ValueError):
            agent.yorumlari_oku([APPLE], [])


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_denetim_ve_gruplar(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "r.xlsx"
            s = agent.calistir([PLAY], [APPLE], cikti, sinir=350, evet=True)
            r, n = s["sonuc"], s["notlar"]
            self.assertEqual(len(r), 21)                                      # yalnız puanlı 3 yorum gönderilmedi
            self.assertEqual(r["GP5"]["duygu"], "olumsuz")                   # 1 yıldız olumlu olamaz
            self.assertEqual(r["GP1"]["duygu"], "olumlu")
            self.assertEqual(r["GP7"]["alinti"], "")                         # uydurma alıntı silindi
            self.assertEqual(r["GP3"]["alinti"], "Plus aldım ama reklamlar hâlâ çıkıyor, para verdim neden reklam görüyorum?")
            self.assertTrue(any("400 karakter" in x for x in n["GP5"]))
            g = {x["ad"]: x["adet"] for x in s["gruplar"]}
            self.assertEqual(g, {"karanlık mod": 4, "sesli okuma": 2, "aylık ödeme seçeneği": 2})
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Olumsuz", "Özellik Talepleri", "Yorumlar", "Sürümler", "Özet"])
            self.assertEqual(wb["Olumsuz"].max_row - 1, 10)                 # 1–2 yıldızlı 10 yorum
            o = json.loads((Path(tmp) / f"{agent.PAKET_ADI}.ozet.json").read_text(encoding="utf-8"))
            self.assertEqual(o["gunluk"]["2026-09-07"], {"yorum": 2, "ortalama_puan": 1.0, "olumsuz": 2})
            self.assertEqual(o["ozellik_talepleri"][0], {"ad": "karanlık mod", "adet": 4})

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_bilinen_talepler_sonraki_pakete_gider(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp, mock.patch.object(agent, "PAKET", 5):
            agent.calistir([PLAY], [], Path(tmp) / "r.xlsx", evet=True)
            self.assertIn("sesli okuma", re.search(r"<bilinen_talepler>(.*)</bilinen_talepler>", MESAJLAR[-1]).group(1))

    def test_komut_satiri(self):
        self.assertEqual(agent.main(["--cevap-siniri", "10"]), 2)


if __name__ == "__main__":
    unittest.main()
