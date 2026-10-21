import json
import sys
import tempfile
import unittest
from datetime import date, datetime, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
SON = date(2026, 9, 30)


def o(gun_once, issue, tur, kurulum, surum="2.3.1", baslik=""):
    z = datetime(2026, 9, 30, 12) - timedelta(days=gun_once)
    return {"zaman": z, "gun": z.date(), "issue": issue, "tur": tur, "kurulum": kurulum, "surum": surum, "platform": "ANDROID", "cihaz": "", "os": "", "baslik": baslik}


class KuralTesti(unittest.TestCase):
    def test_oncelik_kurallari(self):
        olay = [o(i % 7, "cok", "FATAL", f"k{i}") for i in range(10)]                    # 10 kurulum = eşik → P1
        olay += [o(0, "yeni", "FATAL", f"y{i}", "2.10.0") for i in range(3)]             # en yeni sürümde ortaya çıktı, 3 kurulum → P1
        olay += [o(9, "eski", "FATAL", "e1", "2.10.0"), o(0, "eski", "FATAL", "e2", "2.10.0")]  # pencereden önce görülmüş → yeni değil → P2
        olay += [o(0, "anr", "ANR", "a1")]                                                # tek kurulum ANR → P2
        olay += [o(1, "sicra", "NON_FATAL", "s0")] + [o(0, "sicra", "NON_FATAL", f"s{i}") for i in range(5)]   # son gün 5, önceki ort 1/6 → sıçrama → P2
        olay += [o(0, "az", "NON_FATAL", "z1")]                                           # P3
        r = main.analiz(olay, SON, 7, 10)
        p = {s["issue"]: s["oncelik"] for s in r["sorunlar"]}
        self.assertEqual(p, {"cok": "P1", "yeni": "P1", "eski": "P2", "anr": "P2", "sicra": "P2", "az": "P3"})
        self.assertEqual(r["en_yeni"], "2.10.0")                                          # 2.10 > 2.9 > 2.3.1 sayısal sıralama
        s = {s["issue"]: s for s in r["sorunlar"]}
        self.assertTrue(s["sicra"]["sicrama"] and not s["sicra"]["yeni"])
        self.assertEqual((s["eski"]["olay"], s["eski"]["kurulum"]), (1, 1))              # pencere dışındaki olay sayılmaz
        self.assertIn("en az P2", s["anr"]["neden"])

    def test_surum_siralama_ve_cokmesiz(self):
        self.assertGreater(main.surum_anahtari("2.10.0"), main.surum_anahtari("2.9.9"))
        r = main.analiz([o(0, "x", "FATAL", "k1"), o(0, "x", "FATAL", "k1"), o(0, "y", "NON_FATAL", "k2")], SON, 7, 10, {"2.3.1": 200})
        self.assertEqual(r["surumler"]["2.3.1"]["cokme_kurulum"], 1)
        self.assertEqual(str(r["surumler"]["2.3.1"]["cokmesiz"]), "0.995")

    def test_bos_pencere(self):
        with self.assertRaises(ValueError):
            main.analiz([o(30, "x", "FATAL", "k")], SON, 7, 10)


class OrnekVeriTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "c.xlsx"
            r = main.calistir([ORNEK / "crashlytics_olaylari.csv"], cikti, None, 7, 10, ORNEK / "surum_kullanicilari.csv")
            self.assertEqual((r["bas"], r["bitis"], r["en_yeni"]), (date(2026, 9, 24), SON, "2.4.0"))
            p1 = [s["issue"] for s in r["sorunlar"] if s["oncelik"] == "P1"]
            self.assertEqual(p1, ["a1f3c9e2b7d4", "b5e1f7a9d3c6"])                      # yaygın çökme + 2.4.0 ile gelen ödeme çökmesi
            s = {x["issue"]: x for x in r["sorunlar"]}
            self.assertEqual((s["e9a4b6c2d8f1"]["oncelik"], s["e9a4b6c2d8f1"]["sicrama"]), ("P2", True))
            self.assertEqual(sum(x["olay"] for x in r["sorunlar"]), sum(sum(v.values()) - v["etkilenen_kurulum"] for v in r["gunluk"].values()))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Kritik", "Sorunlar", "Sürümler", "Günlük", "Kurallar", "Uyarılar"])
            b = (Path(tmp) / "cokme_bildirimi.txt").read_text(encoding="utf-8")
            self.assertIn("OdemeAkisi.satinAlmaTamamla", b)
            oz = json.loads((Path(tmp) / f"{main.PAKET}.ozet.json").read_text(encoding="utf-8"))
            self.assertEqual((len(oz["gunluk"]), len(oz["kritik"]), oz["en_yeni_surum"]), (7, 2, "2.4.0"))

    def test_tekrar_eden_olay_sayilmaz(self):
        u = []
        olaylar = main.oku([ORNEK / "crashlytics_olaylari.csv"], u)
        self.assertEqual(len(olaylar), len({x["zaman"].isoformat() + x["kurulum"] + x["issue"] for x in olaylar}))
        self.assertEqual(len(olaylar), 227)

    def test_komut_satiri(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(main.main(["--gun", "1"]), 2)


if __name__ == "__main__":
    unittest.main()
