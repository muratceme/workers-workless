import json
import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
FIYAT = {"para": "USD", "birim": D(1000000), "modeller": {"model-a": {"girdi": D(3), "cikti": D(15), "onbellek_okuma": D("0.3")}, "model-a-mini": {"girdi": D(1), "cikti": D(5)}}}


def gunluk_yaz(tmp, satirlar, bas="Zaman,Model,Girdi Token,Çıktı Token,Önbellek Okuma Token,Kullanıcı,Özellik"):
    y = Path(tmp) / "g.csv"
    y.write_text(bas + "\n" + "\n".join(satirlar) + "\n", encoding="utf-8")
    return y


class MaliyetTesti(unittest.TestCase):
    def test_fiyat_ve_onbellek(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = gunluk_yaz(tmp, ["2026-09-01 10:00:00,model-a,1800,650,1200,u1,fal"])
            self.assertEqual(main.oku([y], FIYAT, False, [])[0]["maliyet"], D("0.01551"))         # (1800×3 + 650×15 + 1200×0,3) ÷ 1e6
            self.assertEqual(main.oku([y], FIYAT, True, [])[0]["maliyet"], D("0.01191"))          # girdi önbelleği içeriyor: (600×3 + 650×15 + 1200×0,3) ÷ 1e6

    def test_model_eslestirme_ve_tanimsiz(self):
        self.assertEqual(main.fiyat_bul(FIYAT, "model-a-mini-2026-09")[0], "model-a-mini")     # en uzun ön ek
        self.assertEqual(main.fiyat_bul(FIYAT, "model-a-2026-09")[0], "model-a")
        with tempfile.TemporaryDirectory() as tmp:
            u = []
            c = main.oku([gunluk_yaz(tmp, ["2026-09-01 10:00:00,baska-model,100,100,0,u1,fal"])], FIYAT, False, u)
            self.assertEqual((c[0]["maliyet"], u[0]["onem"], u[0]["kim"]), (D(0), "Yüksek", "baska-model"))

    def test_gunlukteki_maliyet_onceliklidir(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = gunluk_yaz(tmp, ["2026-09-01 10:00:00,model-a,1800,650,0.0123"], "Zaman,Model,Girdi Token,Çıktı Token,Maliyet")
            c = main.oku([y], FIYAT, False, [])[0]
            self.assertEqual((c["maliyet"], c["kaynak"]), (D("0.0123"), "günlük"))


class AnalizTesti(unittest.TestCase):
    def cagri(self, gun, kullanici, maliyet):
        from datetime import datetime
        return {"zaman": datetime(2026, 9, gun, 12), "gun": date(2026, 9, gun), "model": "m", "girdi": 1, "cikti": 1, "okuma": 0, "yazma": 0,
                "kullanici": kullanici, "ozellik": "f", "maliyet": D(maliyet), "kaynak": "fiyat", "hata": False}

    def test_limitler_tahmin_ve_kullanici_basina(self):
        c = [self.cagri(1, "u1", "1")] * 3 + [self.cagri(2, "u2", "1"), self.cagri(10, "u1", "4")]
        r = main.analiz(c, {"aylik_butce": 20, "kullanici_gunluk_istek": 2, "gunluk_butce": 3.5})
        self.assertEqual(r["tahmin"], D(8) / 10 * 30)                                            # 8 USD ÷ 10 gün × 30 gün = 24
        self.assertEqual((r["toplam"], r["kullanici_sayisi"], r["kullanici_basina"]), (D(8), 2, D(4)))
        self.assertEqual([(a["gun"], a["kullanici"]) for a in r["asimlar"]], [(date(2026, 9, 1), "u1")])
        tur = sorted(u["tur"] for u in r["uyarilar"])
        self.assertEqual(tur, ["Aylık bütçe aşılacak", "Günlük bütçe aşıldı", "Kullanıcı limiti aşımı"])   # önceki 7 gün ortalaması 0 → sıçrama sayılmaz

    def test_sicrama_bos_gunleri_sifir_sayar(self):
        c = [self.cagri(g, "u1", "1") for g in (1, 2, 3, 4, 5, 6, 7, 8)] + [self.cagri(9, "u1", "1")] + [self.cagri(9, "u2", "1")]
        self.assertIn("Maliyet sıçraması", [u["tur"] for u in main.analiz(c)["uyarilar"]])            # 2 ≥ 2 × 1
        c = [self.cagri(1, "u1", "1"), self.cagri(8, "u1", "1"), self.cagri(8, "u2", "1")]               # 2–7 boş = 0: ortalama 1/7
        self.assertIn("Maliyet sıçraması", [u["tur"] for u in main.analiz(c)["uyarilar"]])


class OrnekVeriTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "a.xlsx"
            r = main.calistir([ORNEK / "api_cagrilari.csv"], cikti, ORNEK / "fiyatlar.json", ORNEK / "limitler.json", D("41.20"))
            self.assertEqual((r["istek"], r["ilk"], r["son"], r["para"]), (1449, date(2026, 9, 1), date(2026, 9, 30), "USD"))
            self.assertEqual({a["kullanici"] for a in r["asimlar"]}, {"u0077"})
            self.assertIn("model-c-deneme", [u["kim"] for u in r["uyarilar"] if u["tur"] == "Fiyatı tanımsız model"])
            self.assertEqual(sum(1 for c in r["gunluk"].values() for _ in [0]), 30)
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Günlük", "Modeller", "Özellikler", "Kullanıcılar", "Limit Aşımları", "Uyarılar"])
            o = json.loads((Path(tmp) / f"{main.PAKET}.ozet.json").read_text(encoding="utf-8"))
            self.assertEqual((o["para_birimi"], o["kur"], len(o["gunluk"])), ("USD", 41.2, 30))
            self.assertAlmostEqual(sum(g["maliyet"] for g in o["gunluk"].values()), float(r["toplam"]), places=6)

    def test_komut_satiri(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "a.xlsx")]), 0)
            self.assertEqual(main.main(["--kur", "abc"]), 2)


if __name__ == "__main__":
    unittest.main()
