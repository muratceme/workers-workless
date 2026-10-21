import json
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def yaz(klasor, paket, gunluk, **ek):
    y = Path(klasor) / f"{paket}.ozet.json"
    y.write_text(json.dumps({"paket": paket, "surum": 1, "uretildi": ek.pop("uretildi", "2026-10-01T09:00:00"), "gunluk": gunluk, **ek}), encoding="utf-8")
    return y


class DonemTesti(unittest.TestCase):
    def test_araliklar(self):
        t = date(2026, 3, 31)
        self.assertEqual(main.araliklar(t, "gunluk"), ((t, t), (date(2026, 3, 30), date(2026, 3, 30))))
        self.assertEqual(main.araliklar(t, "haftalik"), ((date(2026, 3, 25), t), (date(2026, 3, 18), date(2026, 3, 24))))
        self.assertEqual(main.araliklar(t, "aylik"), ((date(2026, 3, 1), t), (date(2026, 2, 1), date(2026, 2, 28))))   # şubat 28 gün

    def test_toplam_ortalama_agirlikli(self):
        g = {date(2026, 9, 1): {"dau": 10, "yorum": 1, "ortalama_puan": 5}, date(2026, 9, 2): {"dau": 20, "yorum": 3, "ortalama_puan": 1}}
        a = (date(2026, 9, 1), date(2026, 9, 3))
        self.assertEqual((main.deger(g, a, "dau", "t"), main.deger(g, a, "dau", "o"), main.deger(g, a, "ortalama_puan", "a")), (30, 15, 2))
        self.assertEqual(main.kapsam(g, a), (2, 3))
        self.assertIsNone(main.deger(g, (date(2026, 8, 1), date(2026, 8, 2)), "dau", "t"))


class AnalizTesti(unittest.TestCase):
    def test_dikkat_ve_birlesik(self):
        with tempfile.TemporaryDirectory() as tmp:
            yaz(tmp, "uygulama-abonelik-gelir-raporu", {"2026-09-01": {"net": 1000}, "2026-09-02": {"net": 700}}, para_birimi="TRY")
            yaz(tmp, "yapay-zeka-api-maliyet-raporu", {"2026-09-01": {"maliyet": 1.0}, "2026-09-02": {"maliyet": 1.0}}, para_birimi="USD", kur=40,
                ay_sonu_tahmini=40, aylik_butce=30)
            yaz(tmp, "uygulama-kullanici-ve-tutma-analizi", {"2026-09-01": {"dau": 100}, "2026-09-02": {"dau": 50}})
            yaz(tmp, "uygulama-cokme-onceliklendirme", {"2026-09-01": {"fatal_olay": 1}}, kritik=[{"baslik": "X", "tur": "FATAL", "kurulum": 9, "olay": 9, "oncelik": "P1"}])
            k, _ = main.oku([tmp])
            r = main.analiz(k, "gunluk")
            s = {x["ad"]: x for x in r["satirlar"]}
            self.assertAlmostEqual(s["Net gelir (TRY)"]["degisim"], -0.3)
            self.assertAlmostEqual(s["Maliyet / aktif kullanıcı-gün (USD)"]["bu"], 0.02)           # 1 USD ÷ 50 kullanıcı
            self.assertAlmostEqual(s["API maliyeti / net gelir"]["bu"], 40 / 700)                 # 1 USD × 40 ÷ 700 TL
            tur = [d["tur"] for d in r["dikkat"]]
            for beklenen in ("Gelir: Net gelir", "Kullanıcı: Günlük aktif kullanıcı (ort.)", "Yapay zekâ bütçesi", "Çökme P1", "Çökme: eski veri"):
                self.assertIn(beklenen, tur)
            self.assertNotIn("Yapay zekâ: API maliyeti", tur)                                    # değişmedi
            self.assertEqual([x["bolum"] for x in r["satirlar"]].count("Yapay zekâ"), 4)
            bolumler = [x["bolum"] for x in r["satirlar"]]
            self.assertEqual(bolumler, sorted(bolumler, key=bolumler.index))                     # bölümler bölünmeden sıralı

    def test_en_yeni_dosya_ve_taninmayan(self):
        with tempfile.TemporaryDirectory() as tmp:
            (Path(tmp) / "a").mkdir()
            yaz(Path(tmp) / "a", "uygulama-destek-epostasi", {"2026-09-01": {"talep": 1}}, uretildi="2026-09-01T09:00:00")
            yaz(tmp, "uygulama-destek-epostasi", {"2026-09-01": {"talep": 5}}, uretildi="2026-09-02T09:00:00")
            yaz(tmp, "baska-paket", {"2026-09-01": {}})
            k, u = main.oku([tmp])
            self.assertEqual(k["uygulama-destek-epostasi"]["gunluk"][date(2026, 9, 1)]["talep"], 5)
            self.assertEqual(sorted(x["tur"] for x in u), ["Aynı paketten birden çok dosya", "Tanınmayan dosya"])
            with self.assertRaises(ValueError):
                main.oku([Path(tmp) / "yok"])


class OrnekVeriTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "y.xlsx"
            r = main.calistir([ORNEK], cikti, "aylik")
            self.assertEqual((r["t"], len(r["kaynak"])), (date(2026, 9, 30), 6))
            s = {x["ad"]: x for x in r["satirlar"]}
            self.assertAlmostEqual(s["Net gelir (TRY)"]["bu"], 37657.92, places=2)               # gelir paketinin eylül toplamı
            self.assertIn("API maliyeti / net gelir", r["metin"])                                 # kur yok → neden hesaplanmadığı yazılır
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Dikkat", "Günlük", "Kaynaklar", "Uyarılar"])
            self.assertTrue(cikti.with_suffix(".md").read_text(encoding="utf-8").startswith("# Uygulama Yönetim Özeti · Aylık"))

    def test_komut_satiri(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "y.xlsx"), "--donem", "haftalik"]), 0)
            self.assertEqual(main.main(["--tarih", "2026-09-30"]), 2)
            self.assertEqual(main.main(["--kaynak", tmp]), 1)                                     # boş klasör


if __name__ == "__main__":
    unittest.main()
