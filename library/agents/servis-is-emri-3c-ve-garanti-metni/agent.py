"""
Servis İş Emri 3C ve Garanti Metni — Gerisi Bende AI Agent
Otomotiv Servis › Servis Danışmanı (Garanti Uzmanı da kullanır)

Müşteri şikâyeti ve teknisyen notlarından iş emrine ve garanti başvurusuna girecek 3C metnini hazırlar:
Şikâyet (Complaint), Neden (Cause), Çözüm (Correction); değişen parçalar ve işçilik.
1. Kod iş emri dosyasını okur ("İş Emri: … | Plaka | Model | Trafiğe çıkış | KM | Tarih" satırı, ardından "Müşteri:" ve
   "Teknisyen:" satırları); telefon, e-posta, kimlik numarası maskelenir. Garanti koşulları tablosu okunur.
2. Model 3C metnini yazar, değişen parçaları (parça no, ad, adet) ve işçilik saatini çıkarır, garanti başvurusu önerisini
   gerekçesiyle verir.
3. Kod denetler: parça no ve işçilik saati teknisyen notunda aynen geçiyor mu; şikâyet müşterinin sözüne, neden ve çözüm
   teknisyen notuna dayanıyor mu (notta olmayan sayı yok); şikâyet bölümüne teknisyen bulgusu karışmış mı.
4. Kod garanti kapsamını hesaplar (araç yaşı ve km, modelin garanti yılı / km sınırı) ve garanti dışı olabilecek işaretleri
   (darbe, kaza, yanlış yakıt, modifiye, sıvı girişi, bakım yapılmamış) bulur; bunlar varken garanti önerilmişse işaretler.

Kullanım:
    python agent.py                                                         # örnek 2 iş emri
    python agent.py --girdi is_emirleri.txt --garanti garanti_kosullari.csv
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
RENK = {"Garanti kapsamında": "E3F5E1", "Kontrol": "FFF4CE", "Garanti dışı": "FDE2E1", "Ücretli": "E8F0FE"}
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}
BASLIK = re.compile(r"^\s*İş Emri\s*:\s*(.+)$", re.I)
DISI = re.compile(r"\b(darbe|carpma|kaza|yanlis yakit|modifiye|sivi girisi|su girisi|su almis|bakim yapilmamis|bakimi yapilmamis|"
                  r"kullanici hatasi|asiri yuk|yangin|sel)\w*")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def para(x) -> Decimal | None:
    s = re.sub(r"[^\d,.\-]", "", str(x or ""))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1 or (s.count(".") == 1 and len(s.split(".")[1]) == 3):
        s = s.replace(".", "")
    try:
        return Decimal(s) if s not in ("", "-", ".") else None
    except InvalidOperation:
        return None


def sayilar(metin: str) -> set[Decimal]:
    temiz = re.sub(r"\b\d{1,2}\.\d{1,2}\.\d{4}\b|\b[A-Z0-9]+-[A-Z0-9-]+\b", " ", str(metin))
    return {v for x in re.findall(r"\d+(?:[.,]\d+)*", temiz) if (v := para(x)) is not None}


def tarih_coz(x) -> date | None:
    m = re.search(r"\b(\d{1,2})\.(\d{1,2})\.(\d{4})\b", str(x or ""))
    if not m:
        return None
    try:
        return date(int(m[3]), int(m[2]), int(m[1]))
    except ValueError:
        return None


# ----------------------------------------------------------------------------
# Girdi
# ----------------------------------------------------------------------------

def is_emirlerini_oku(yol: Path) -> list[dict]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: iş emri dosyası bulunamadı.")
    out, son = [], None
    for s in belge.metin_oku(yol).splitlines():
        if not s.strip():
            continue
        m = BASLIK.match(s)
        if m:
            alan = [x.strip() for x in m[1].split("|")]
            bilgi = {"no": alan[0]}
            for a in alan[1:]:
                k, _, v = a.partition(":")
                bilgi[katla(k)] = v.strip()
            out.append({"no": bilgi["no"], "plaka": bilgi.get("plaka", ""), "model": bilgi.get("model", ""), "cikis": tarih_coz(bilgi.get("trafige cikis")),
                        "km": para(bilgi.get("km")), "tarih": tarih_coz(bilgi.get("tarih")), "musteri": "", "teknisyen": ""})
            son = None
            continue
        if not out:
            continue
        k = re.match(r"^\s*(Müşteri|Musteri|Teknisyen)\s*:\s*(.*)$", s, re.I)
        if k:
            son = "musteri" if katla(k[1]) == "musteri" else "teknisyen"
            out[-1][son] = llm.maskele(k[2].strip())
        elif son:
            out[-1][son] += " " + llm.maskele(s.strip())
    out = [x for x in out if x["musteri"] or x["teknisyen"]]
    if not out:
        raise llm.LLMHatasi(f"{yol.name}: 'İş Emri: …' başlıklı kayıt yok.")
    return out


def garanti_oku(yol: Path | None) -> dict[str, tuple[int, Decimal]]:
    if yol is None:
        return {}
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: garanti koşulları dosyası bulunamadı.")
    metin = yol.read_text(encoding="utf-8-sig")
    s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=metin.split("\n")[0].count)))
    return {katla(r[0]): (int(para(r[1]) or 0), para(r[2]) or Decimal(0)) for r in s[1:] if len(r) >= 3 and r[0].strip()}


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

PARCA = {"type": "object", "properties": {"parca_no": S, "ad": S, "adet": S}, "required": ["parca_no", "ad", "adet"], "additionalProperties": False}
SEMA = {"type": "object", "properties": {
    "sikayet": S, "neden": S, "cozum": S, "parcalar": {"type": "array", "items": PARCA}, "iscilik_saat": S,
    "garanti_onerisi": {"type": "string", "enum": ["Garanti başvurusu", "Ücretli", "Karar için bilgi gerekli"]}, "gerekce": S,
    "eksik_bilgiler": {"type": "array", "items": S}},
    "required": ["sikayet", "neden", "cozum", "parcalar", "iscilik_saat", "garanti_onerisi", "gerekce", "eksik_bilgiler"], "additionalProperties": False}


def denetle(ie: dict, c: dict) -> list[str]:
    on, sorun = ie["no"], []
    tek, mus = ie["teknisyen"], ie["musteri"]
    for p in c["parcalar"]:
        if p["parca_no"].strip() and p["parca_no"].strip() not in tek:
            sorun.append(f"{on}: parça no '{p['parca_no']}' teknisyen notunda aynen geçmiyor")
        if not p["parca_no"].strip():
            sorun.append(f"{on}: '{p['ad']}' için parça no yok; notta yoksa eksik bilgilere yazın")
    if c["iscilik_saat"].strip() and (para(c["iscilik_saat"]) is None or para(c["iscilik_saat"]) not in sayilar(tek)):
        sorun.append(f"{on}: işçilik '{c['iscilik_saat']}' teknisyen notunda yazmıyor")
    fazla = sorted(sayilar(c["sikayet"]) - sayilar(mus), key=float)
    if fazla:
        sorun.append(f"{on}: şikâyette müşterinin söylemediği sayı: {', '.join(f'{x:g}' for x in fazla)}")
    fazla = sorted(sayilar(c["neden"] + " " + c["cozum"]) - sayilar(tek), key=float)
    if fazla:
        sorun.append(f"{on}: neden / çözümde teknisyen notunda olmayan sayı: {', '.join(f'{x:g}' for x in fazla)}")
    teknik = [k for k in re.findall(r"[a-zçğıöşü]{5,}", katla(tek)) if k not in katla(mus)]
    sizan = [k for k in ("gergi", "conta", "sensor", "kablo", "zincir", "basinc", "arizali", "degistir", "kacak", "yatak", "rulman", "enjektor")
             if k in katla(c["sikayet"]) and k not in katla(mus) and any(k in w for w in teknik)]
    if sizan:
        sorun.append(f"{on}: şikâyet bölümünde teknisyen bulgusu var ({', '.join(sizan)}); şikâyet yalnız müşterinin anlattığı belirtidir")
    for a, ad in (("sikayet", "Şikâyet"), ("neden", "Neden"), ("cozum", "Çözüm")):
        if not c[a].strip():
            sorun.append(f"{on}: {ad} boş")
    return sorun


def kapsam(ie: dict, c: dict, kosullar: dict[str, tuple[int, Decimal]]) -> tuple[str, list[str]]:
    """Kod: garanti süresi / km kontrolü ve garanti dışı işaretleri."""
    b = []
    k = kosullar.get(katla(ie["model"]))
    if k and ie["cikis"] and ie["tarih"]:
        yil, km = k
        try:
            bitis = ie["cikis"].replace(year=ie["cikis"].year + yil)
        except ValueError:                                  # 29 Şubat
            bitis = date(ie["cikis"].year + yil, 2, 28)
        if ie["tarih"] > bitis:
            b.append(f"Garanti süresi dolmuş ({yil} yıl: {bitis:%d.%m.%Y})")
        if ie["km"] is not None and km and ie["km"] > km:
            b.append(f"Garanti km sınırı aşılmış ({ie['km']:,.0f} > {km:,.0f})".replace(",", "."))
    elif kosullar:
        b.append("Model, trafiğe çıkış tarihi ya da iş emri tarihi eksik: garanti süresi hesaplanamadı")
    disi = sorted({m.group(0) for m in DISI.finditer(katla(ie["teknisyen"]))})
    if disi:
        b.append(f"Garanti dışı olabilecek işaret: {', '.join(disi)}")
    if any(x.startswith("Garanti süresi dolmuş") or x.startswith("Garanti km") for x in b):
        durum = "Garanti dışı"
    elif c["garanti_onerisi"] == "Garanti başvurusu":
        durum = "Kontrol" if disi or any("hesaplanamadı" in x for x in b) else "Garanti kapsamında"
    elif c["garanti_onerisi"] == "Ücretli":
        durum = "Ücretli"
    else:
        durum = "Kontrol"
    return durum, b


def uret(emirler: list[dict], duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    sonuc = dict(onceki or {})
    hedef = emirler if duzeltme is None else [e for e in emirler if any(s.startswith(e["no"] + ":") for s in duzeltme)]
    for e in hedef:
        km = f"{e['km']:,.0f}".replace(",", ".") if e["km"] is not None else "-"
        mesaj = (f"<is_emri no=\"{e['no']}\" model=\"{e['model']}\" km=\"{km}\">\n<musteri>{e['musteri']}</musteri>\n"
                 f"<teknisyen>{e['teknisyen']}</teknisyen>\n</is_emri>")
        d = [s.split(": ", 1)[1] for s in (duzeltme or []) if s.startswith(e["no"] + ":")]
        if d:
            mesaj += ("\n<duzeltme>\nÖnceki metin: " + json.dumps(sonuc.get(e["no"]), ensure_ascii=False) + "\nKodun bulduğu sorunlar:\n"
                      + "\n".join(f"- {s}" for s in d) + "\n</duzeltme>")
        sonuc[e["no"]] = llm.json_iste(sistem, mesaj, SEMA)
    return sonuc


def duzeltme_dongusu(uret_f, denetle_f, tur: int) -> tuple[dict, list[str], int]:
    taslak = uret_f(None, None)
    kalan = denetle_f(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret_f(kalan, taslak)
        kalan = denetle_f(taslak)
    return taslak, kalan, yapilan


def rapor_yaz(cikti: Path, emirler: list[dict], sonuc: dict, kalan: list[str], tur: int) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "İş Emirleri"
    bas = ["İş Emri", "Plaka", "Model", "KM", "Şikâyet (Complaint)", "Neden (Cause)", "Çözüm (Correction)", "Değişen Parçalar", "İşçilik (saat)",
           "Model Önerisi", "Gerekçe", "Garanti Durumu (kod)", "Kod Bulguları", "Eksik Bilgiler", "Son Kontrol"]
    ws.append(bas)
    for h in ws[1]:
        h.fill, h.font, h.alignment = BASLIK_DOLGU, BASLIK_YAZI, UST
    for j, w in enumerate((14, 11, 14, 9, 40, 45, 45, 40, 8, 16, 40, 16, 45, 35, 30), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for e in emirler:
        c = sonuc[e["no"]]
        sor = [x.split(": ", 1)[1] for x in kalan if x.startswith(e["no"] + ":")]
        ws.append([e["no"], e["plaka"], e["model"], None if e["km"] is None else float(e["km"]), c["sikayet"], c["neden"], c["cozum"],
                   "\n".join(f"{p['parca_no']} · {p['ad']} · {p['adet'] or 1} adet" for p in c["parcalar"]), None if para(c["iscilik_saat"]) is None else float(para(c["iscilik_saat"])),
                   c["garanti_onerisi"], c["gerekce"], e["durum"], "\n".join(e["bulgu"]), "\n".join(f"• {x}" for x in c["eksik_bilgiler"]), "; ".join(sor) or "Geçti"])
        n = ws.max_row
        for j in (5, 6, 7, 8, 10, 11):
            ws.cell(n, j).fill = AI_DOLGU
        ws.cell(n, 12).fill = PatternFill("solid", fgColor=RENK[e["durum"]])
        ws.cell(n, 15).fill = KIRMIZI if sor else YESIL
        for h in ws[n]:
            h.alignment = UST
    k = wb.create_sheet("Kontrol")
    k.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    k.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    k.cell(2, 1).fill = KIRMIZI if kalan else YESIL
    k.append(["Önemli: 3C metni ve garanti önerisi yapay zekâ taslağıdır (mor). Garanti onayı üreticinin / distribütörün garanti politikasına "
              "göre garanti uzmanınca verilir. Garanti süresi / km kontrolü verdiğiniz garanti koşulları tablosuna göre kodla yapılmıştır."])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(girdi: Path, cikti: Path, garanti_yolu: Path | None = None, tur: int = 2, evet: bool = False) -> dict:
    emirler = is_emirlerini_oku(girdi)
    kosullar = garanti_oku(garanti_yolu)
    print(f"[OK] {len(emirler)} iş emri · garanti koşulu: {len(kosullar)} model")
    llm.onay_al(f"{len(emirler)} iş emrinin müşteri şikâyeti ve teknisyen notları gönderilecek (model ve km dahil; plaka gönderilmez; telefon, "
                "e-posta, TCKN maskeli).", evet)
    sonuc, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(emirler, d, o), lambda x: [s for e in emirler for s in denetle(e, x[e["no"]])], tur)
    for e in emirler:
        e["durum"], e["bulgu"] = kapsam(e, sonuc[e["no"]], kosullar)
    rapor_yaz(cikti, emirler, sonuc, kalan, yapilan)
    return {"emirler": emirler, "sonuc": sonuc, "kalan": kalan, "tur": yapilan}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Müşteri şikâyeti ve teknisyen notlarından iş emri / garanti 3C metni hazırlar; garanti kapsamını kodla kontrol eder.")
    p.add_argument("--girdi", type=Path, default=ORNEK / "is_emirleri.txt", help="İş emirleri (.txt/.docx)")
    p.add_argument("--garanti", type=Path, default=None, help="Garanti koşulları (.csv: Model; Garanti Yıl; Garanti KM)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "is_emri_3c.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    garanti = a.garanti or (ORNEK / "garanti_kosullari.csv" if a.girdi == ORNEK / "is_emirleri.txt" else None)
    try:
        s = calistir(a.girdi, a.cikti, garanti, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    for e in s["emirler"]:
        print(f"[OK] {e['no']} · {s['sonuc'][e['no']]['garanti_onerisi']} → {e['durum']}")
        for b in e["bulgu"]:
            print(f"[!]   {b}")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
