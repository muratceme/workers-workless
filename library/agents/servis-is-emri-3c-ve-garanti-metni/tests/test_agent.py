"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import json
import os
import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
IE = ORNEK / "is_emirleri.txt"
GAR = ORNEK / "garanti_kosullari.csv"
CEVAP = json.loads((Path(__file__).parent / "sahte_cevaplar.json").read_text(encoding="utf-8"))
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    """IE-2026-1023'te ilk turda şikâyet bölümüne teknisyen bulgusunu ('zincir gergisi kaynaklı') ekler."""
    MESAJLAR.append(kullanici)
    if "IE-2026-1031" in kullanici:
        return copy.deepcopy(CEVAP["1031"])
    return copy.deepcopy(CEVAP["1023" if "<duzeltme>" in kullanici else "1023_hatali"])


def emirler():
    return {e["no"]: e for e in agent.is_emirlerini_oku(IE)}


class KodTesti(unittest.TestCase):
    def test_okuma(self):
        e = emirler()
        x = e["IE-2026-1023"]
        self.assertEqual((x["plaka"], x["model"], x["cikis"], x["km"], x["tarih"]), ("34 ABC 123", "Kurgu X1 1.5", date(2024, 3, 12), D(48200), date(2026, 10, 15)))
        self.assertTrue(x["musteri"].startswith("Sabah soğukta") and "13540-KX100" in x["teknisyen"])
        self.assertEqual(agent.garanti_oku(GAR)["kurgu x1 1 5"], (3, D(100000)))

    def test_denetim(self):
        e = emirler()
        self.assertEqual([agent.denetle(e["IE-2026-1023"], CEVAP["1023"]), agent.denetle(e["IE-2026-1031"], CEVAP["1031"])], [[], []])
        self.assertEqual(agent.denetle(e["IE-2026-1023"], CEVAP["1023_hatali"]),
                         ["IE-2026-1023: şikâyet bölümünde teknisyen bulgusu var (gergi, zincir); şikâyet yalnız müşterinin anlattığı belirtidir"])
        c = copy.deepcopy(CEVAP["1023"])
        c["parcalar"][0]["parca_no"] = "13540-KX101"
        c["parcalar"].append({"parca_no": "", "ad": "Yağ", "adet": "4"})
        c["iscilik_saat"] = "3"
        c["neden"] += " 5 dakikada ısınıyor."
        s = agent.denetle(e["IE-2026-1023"], c)
        for parca in ("13540-KX101", "'Yağ' için parça no yok", "işçilik '3'", "olmayan sayı: 5"):
            self.assertTrue(any(parca in x for x in s), parca)

    def test_kapsam(self):
        e = emirler()
        k = agent.garanti_oku(GAR)
        self.assertEqual(agent.kapsam(e["IE-2026-1023"], CEVAP["1023"], k), ("Garanti kapsamında", []))
        self.assertEqual(agent.kapsam(e["IE-2026-1031"], CEVAP["1031"], k), ("Ücretli", ["Garanti dışı olabilecek işaret: darbe"]))
        self.assertEqual(agent.kapsam(e["IE-2026-1031"], CEVAP["1031"] | {"garanti_onerisi": "Garanti başvurusu"}, k)[0], "Kontrol")
        eski = e["IE-2026-1023"] | {"cikis": date(2023, 10, 1), "km": D(120000)}
        d, b = agent.kapsam(eski, CEVAP["1023"], k)
        self.assertEqual((d, len(b)), ("Garanti dışı", 2))
        self.assertIn("01.10.2026", b[0])
        artik = e["IE-2026-1023"] | {"cikis": date(2024, 2, 29), "tarih": date(2027, 3, 1)}
        self.assertEqual(agent.kapsam(artik, CEVAP["1023"], k)[0], "Garanti dışı")         # 29 Şubat + 3 yıl = 28.02.2027
        self.assertIn("hesaplanamadı", agent.kapsam(e["IE-2026-1023"] | {"model": "Bilinmeyen"}, CEVAP["1023"], k)[1][0])


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ornek_ve_duzeltme(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "i.xlsx"
            self.assertEqual(agent.main(["--evet", "--cikti", str(cikti)]), 0)
            self.assertEqual(len(MESAJLAR), 3)
            for m in MESAJLAR:
                self.assertNotIn("34 ABC 123", m)                          # plaka gönderilmez
            ws = load_workbook(cikti)["İş Emirleri"]
            self.assertEqual([ws.cell(2, j).value for j in (1, 9, 10, 12)], ["IE-2026-1023", 2.5, "Garanti başvurusu", "Garanti kapsamında"])
            self.assertEqual([ws.cell(3, j).value for j in (10, 12)], ["Ücretli", "Ücretli"])
            self.assertNotIn("gergi", ws.cell(2, 5).value)

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(agent.main(["--girdi", str(Path(tmp) / "yok.txt"), "--evet"]), 1)
            self.assertEqual(agent.main(["--garanti", str(Path(tmp) / "yok.csv"), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
