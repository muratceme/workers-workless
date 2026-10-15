"""
İade Nedeni Analizi — Workers / Workless AI Agent
E-ticaret Operasyon ve Kargo › E-ticaret Operasyon Uzmanı

1. Kod iade açıklamalarını okur; boş açıklamalar "Belirtilmemiş" sayılır ve modele gönderilmez. Açık ifadeler için kural
   işaretleri çıkarır ("küçük / dar geldi" → beden küçük, "büyük / bol geldi" → beden büyük, "yanlış ürün / farklı beden
   geldi" → yanlış gönderim, "bozuk / sökük / yırtık / lekeli" → kusurlu, "geç geldi" → geç teslimat).
2. Model her iadeyi neden kategorisine ayırır ve dayandığı ifadeyi alıntılar.
3. Kod denetler: her iade sınıflandırılmış mı, alıntı açıklamada var mı, tek bir kural işareti varken model farklı kategori
   vermiş mi. Sorun varsa modele geri gönderilir (en fazla --tur kez).
4. Kod analiz eder: ürün bazında iade oranı (iade / satış), --esik %'yi (varsayılan 15) aşan ürünler, beden yönü
   (küçük / büyük geldi baskınsa kalıp notu önerisi), kusur ve yanlış gönderim (depo / tedarikçi aksiyonu), görsel /
   renk beklenti farkı (ürün görseli kontrolü).

Kullanım:
    python agent.py                                      # örnek iadeler + satış adetleri
    python agent.py --girdi iadeler.xlsx --satislar satislar.xlsx --esik 20
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
PAKET = 40
KUCUK, BUYUK, KALIP, KUSUR, YANLIS = "Beden küçük geldi", "Beden büyük geldi", "Kalıp / duruş beğenilmedi", "Kusurlu / hasarlı ürün", "Yanlış ürün / beden gönderildi"
RENK, KALITE, GEC, VAZGECTI, BELIRSIZ = "Beklenti farkı: renk / görsel", "Beklenti farkı: kumaş / kalite", "Geç teslimat", "Vazgeçti / fikir değişikliği", "Belirtilmemiş"
KATEGORILER = [KUCUK, BUYUK, KALIP, KUSUR, YANLIS, RENK, KALITE, GEC, VAZGECTI, "Diğer", BELIRSIZ]
KURALLAR = {
    KUCUK: re.compile(r"\b(kucuk geldi|dar geldi|dar|kisa geldi|bir beden buyugunu|sikti|sikiyor)\b"),
    BUYUK: re.compile(r"\b(buyuk geldi|bol geldi|bol|uzun geldi|bir beden kucugunu)\b"),
    YANLIS: re.compile(r"(yanlis urun|yanlis beden|yanlis renk|farkli beden geldi|baska urun geldi|siparis ettim \w+ beden geldi)"),
    KUSUR: re.compile(r"\b(bozuk|sokuk|yirtik|lekeli|defolu|kirik|delik|hasarli|kopuk)\b"),
    GEC: re.compile(r"(gec geldi|gecikti|gec teslim|gelmedi zamaninda)"),
}
SEMA = {
    "type": "object",
    "properties": {"iadeler": {"type": "array", "items": {"type": "object", "properties": {"id": S, "kategori": {"type": "string", "enum": KATEGORILER}, "alinti": S},
                                                          "required": ["id", "kategori", "alinti"], "additionalProperties": False}}},
    "required": ["iadeler"],
    "additionalProperties": False,
}
AKSIYON = {
    KUCUK: "Ölçü tablosunu numuneyle karşılaştırın; ürün sayfasına 'kalıbı dar, bir beden büyük alın' notu ekleyin",
    BUYUK: "Ölçü tablosunu numuneyle karşılaştırın; ürün sayfasına 'kalıbı geniş, bir beden küçük alın' notu ekleyin",
    KUSUR: "Stoktaki ürünlerde kalite kontrol yapın; tedarikçiye kusur bildirimi açın",
    YANLIS: "Depoda toplama / paketleme kontrolünü ve ürün-barkod eşleşmesini kontrol edin",
    RENK: "Ürün fotoğraflarının renk doğruluğunu ve ışığını kontrol edin; açıklamaya renk notu ekleyin",
    KALITE: "Ürün açıklamasındaki kumaş / gramaj bilgisini gözden geçirin; tedarikçi numunesiyle karşılaştırın",
    GEC: "Kargo firması teslim sürelerini ve depodan çıkış süresini kontrol edin",
}


def kural(metin: str) -> str | None:
    k = katla(metin)
    eslesen = [ad for ad, d in KURALLAR.items() if d.search(k)]
    if YANLIS in eslesen:
        eslesen = [YANLIS]
    return eslesen[0] if len(eslesen) == 1 else None


def iadeleri_oku(yol: Path) -> list[dict]:
    r = kayitlar(yol, {"id": ("İade No", "İade Kodu", "No"), "siparis": ("Sipariş No", "Sipariş"), "urun": ("Ürün Kodu", "Stok Kodu", "SKU"),
                       "ad": ("Ürün Adı", "Ürün"), "beden": ("Beden", "Varyant"), "tarih": ("İade Tarihi", "Tarih"),
                       "aciklama": ("İade Açıklaması", "Açıklama", "İade Nedeni", "Müşteri Notu")}, ("urun",))
    return [{"id": str(x["id"] or f"IA{i}"), "siparis": str(x["siparis"]), "urun": str(x["urun"]), "ad": str(x["ad"]), "beden": str(x["beden"]),
             "tarih": tarih_coz(x["tarih"]), "aciklama": str(x["aciklama"] or "").strip()} for i, x in enumerate(r, 1) if x["urun"]]


def satislari_oku(yol: Path) -> dict[str, float]:
    r = kayitlar(yol, {"urun": ("Ürün Kodu", "Stok Kodu", "SKU"), "adet": ("Satış Adedi", "Satış", "Adet", "Satılan")}, ("urun", "adet"))
    out = {}
    for x in r:
        try:
            out[str(x["urun"])] = out.get(str(x["urun"]), 0) + float(sayi_norm(x["adet"]))
        except ValueError:
            pass
    return out


def denetle(iade: list[dict], sonuc: dict) -> list[str]:
    sorun = []
    for x in iade:
        if not x["aciklama"]:
            continue
        s = sonuc.get(x["id"])
        if not s:
            sorun.append(f"{x['id']}: sınıflandırılmadı")
            continue
        if s["kategori"] != BELIRSIZ and not metinde_var(s["alinti"], x["aciklama"]):
            sorun.append(f"{x['id']}: alıntı açıklamada yok: '{s['alinti'][:50]}'")
        k = kural(x["aciklama"])
        if k and s["kategori"] != k:
            sorun.append(f"{x['id']}: açıklama ('{x['aciklama'][:60]}') açıkça '{k}' diyor, '{s['kategori']}' verilmiş")
    return sorun


def uret(iade: list[dict], duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    dolu = [x for x in iade if x["aciklama"]]
    sonuc = {}
    if duzeltme is not None:  # düzeltmede yalnız sorunlu iadeler yeniden sorulur
        sorunlu = {s.split(":")[0] for s in duzeltme}
        sonuc.update({a: b for a, b in (onceki or {}).items() if a not in sorunlu})
        dolu = [x for x in dolu if x["id"] in sorunlu]
    for j in range(0, len(dolu), PAKET):
        parca = dolu[j:j + PAKET]
        ids = {x["id"] for x in parca}
        d = [s for s in (duzeltme or []) if s.split(":")[0] in ids]
        if duzeltme is not None and not d:
            sonuc.update({a: b for a, b in (onceki or {}).items() if a in ids})
            continue
        satir = "\n".join(f'<iade id="{x["id"]}" urun="{x["ad"] or x["urun"]}" beden="{x["beden"]}">{llm.maskele(x["aciklama"])}</iade>' for x in parca)
        mesaj = f"<kategoriler>{'; '.join(KATEGORILER)}</kategoriler>\n<iadeler>\n{satir}\n</iadeler>" + duzeltme_blogu(d, [b for a, b in (onceki or {}).items() if a in ids])
        for s in llm.json_iste(sistem, mesaj, SEMA).get("iadeler", []):
            if s.get("id") in ids:
                sonuc[s["id"]] = s
    return sonuc


def kategori(x: dict, sonuc: dict) -> str:
    return sonuc.get(x["id"], {}).get("kategori", BELIRSIZ) if x["aciklama"] else BELIRSIZ


def urun_analizi(iade, sonuc, satis: dict | None, esik: float, beden_min: int) -> list[dict]:
    gr = defaultdict(list)
    for x in iade:
        gr[x["urun"]].append(x)
    out = []
    for u, l in gr.items():
        say = Counter(kategori(x, sonuc) for x in l)
        sat = (satis or {}).get(u)
        oran = round(100 * len(l) / sat, 1) if sat else None
        notlar = []
        if oran is not None and oran > esik:
            notlar.append(f"İade oranı %{oran:g} > %{esik:g}")
        k, b = say[KUCUK], say[BUYUK]
        if k >= beden_min and k >= 2 * b:
            notlar.append(f"Beden küçük ({k} iade): " + AKSIYON[KUCUK])
        elif b >= beden_min and b >= 2 * k:
            notlar.append(f"Beden büyük ({b} iade): " + AKSIYON[BUYUK])
        for kat in (KUSUR, YANLIS, RENK, KALITE, GEC):
            if say[kat]:
                notlar.append(f"{kat} ({say[kat]}): {AKSIYON[kat]}")
        out.append({"urun": u, "ad": next((x["ad"] for x in l if x["ad"]), ""), "iade": len(l), "satis": sat, "oran": oran, "dagilim": say, "notlar": notlar})
    return sorted(out, key=lambda r: (-(r["oran"] if r["oran"] is not None else -1), -r["iade"]))


def rapor_yaz(cikti: Path, iade, sonuc, analiz, esik, kalan, tur) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ürün Analizi"
    kats = [k for k in KATEGORILER if any(r["dagilim"][k] for r in analiz)]
    baslik(ws, ["Ürün Kodu", "Ürün Adı", "Satış", "İade", "İade Oranı %", *kats, "Önerilen Aksiyonlar"], (11, 18, 8, 7, 11, *[11] * len(kats), 80))
    for r in analiz:
        ws.append([r["urun"], r["ad"], r["satis"], r["iade"], r["oran"], *[r["dagilim"][k] or "" for k in kats], "\n".join(r["notlar"])])
        if r["oran"] is not None and r["oran"] > esik:
            boya(ws, [5], KIRMIZI)
    son_not(ws, "Önemli: Kategoriler yapay zekâ sınıflandırmasıdır; açıklama yazılmamış iadeler 'Belirtilmemiş' sayılır.", kalan, tur)
    hizala(ws)
    i = wb.create_sheet("İadeler")
    baslik(i, ["İade No", "Sipariş No", "Ürün Kodu", "Ürün Adı", "Beden", "İade Tarihi", "Açıklama", "Kategori", "Alıntı", "Kural İşareti"],
           (10, 10, 10, 16, 7, 11, 50, 28, 30, 22))
    for x in iade:
        s = sonuc.get(x["id"], {})
        i.append([x["id"], x["siparis"], x["urun"], x["ad"], x["beden"], x["tarih"], x["aciklama"], kategori(x, sonuc), s.get("alinti", ""), kural(x["aciklama"]) or ""])
        i.cell(i.max_row, 6).number_format = "DD.MM.YYYY"
        boya(i, [8, 9], AI_DOLGU)
    i.freeze_panes = "B2"
    i.auto_filter.ref = f"A1:J{i.max_row}"
    hizala(i)
    p = wb.create_sheet("Neden Dağılımı")
    baslik(p, ["Kategori", "İade", "%"], (32, 8, 8))
    say = Counter(kategori(x, sonuc) for x in iade)
    for k, n in say.most_common():
        p.append([k, n, round(100 * n / len(iade), 1)])
    g = BarChart()
    g.type = "bar"
    g.title = "İade nedenleri"
    g.add_data(Reference(p, min_col=2, min_row=1, max_row=p.max_row), titles_from_data=True)
    g.set_categories(Reference(p, min_col=1, min_row=2, max_row=p.max_row))
    g.height, g.width = 8, 16
    p.add_chart(g, "E2")
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(girdi: Path, cikti: Path, satis_yolu: Path | None = None, esik: float = 15, beden_min: int = 2, tur: int = 2, evet: bool = False) -> dict:
    iade = iadeleri_oku(girdi)
    if not iade:
        raise llm.LLMHatasi(f"{girdi.name}: iade yok.")
    satis = satislari_oku(satis_yolu) if satis_yolu else None
    dolu = sum(1 for x in iade if x["aciklama"])
    print(f"[OK] {len(iade)} iade · açıklamalı {dolu} · kural işaretli {sum(1 for x in iade if kural(x['aciklama']))}")
    llm.onay_al(f"{dolu} iade açıklaması, ürün adı ve beden gönderilecek (sipariş no ve müşteri bilgisi gönderilmez).", evet)
    sonuc, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(iade, d, o), lambda x: denetle(iade, x), tur)
    analiz = urun_analizi(iade, sonuc, satis, esik, beden_min)
    rapor_yaz(cikti, iade, sonuc, analiz, esik, kalan, yapilan)
    return {"iadeler": iade, "sonuc": sonuc, "analiz": analiz, "kalan": kalan, "tur": yapilan}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    p = argparse.ArgumentParser(description="İade açıklamalarını neden kategorilerine ayırır, ürün bazında iade oranı ve aksiyon çıkarır.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "iadeler.csv",
                   help="İadeler (.xlsx/.csv): Ürün Kodu, İade Açıklaması; İade No, Sipariş No, Ürün Adı, Beden, İade Tarihi isteğe bağlı")
    p.add_argument("--satislar", type=Path, help="İsteğe bağlı satış adetleri (aynı dönem): Ürün Kodu, Satış Adedi")
    p.add_argument("--esik", type=float, default=15, help="Bu iade oranının (%%) üstündeki ürünleri işaretle (varsayılan 15)")
    p.add_argument("--beden-min", type=int, default=2, help="Beden notu önerisi için en az küçük / büyük geldi iadesi (varsayılan 2)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "iade_analizi.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    if a.girdi == BURASI / "ornek_veri" / "iadeler.csv" and not a.satislar:
        a.satislar = BURASI / "ornek_veri" / "satislar.csv"
    try:
        s = calistir(a.girdi, a.cikti, a.satislar, a.esik, a.beden_min, a.tur, a.evet)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    for r in s["analiz"]:
        if r["notlar"]:
            print(f"[!] {r['urun']} {r['ad']}: " + " | ".join(n.split(":")[0] for n in r["notlar"]))
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
