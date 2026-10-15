"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

IADE = KLASOR / "ornek_veri" / "iadeler.csv"
SATIS = KLASOR / "ornek_veri" / "satislar.csv"
MESAJLAR = []
DOGRU = {"IA-9001": (agent.KUCUK, "küçük geldi"), "IA-9002": (agent.KUCUK, "Bel çok dar"), "IA-9003": (agent.KUCUK, "Küçük geldi"),
         "IA-9004": (agent.RENK, "Rengi fotoğraftakinden çok farklı"), "IA-9005": (agent.KUCUK, "Boyu kısa geldi"), "IA-9006": (agent.KUSUR, "Fermuarı bozuk"),
         "IA-9007": (agent.KUSUR, "Dikiş yeri sökük"), "IA-9008": (agent.VAZGECTI, "Vazgeçtim"), "IA-9009": (agent.YANLIS, "yanlış ürün"),
         "IA-9010": (agent.KALITE, "Kumaşı çok ince"), "IA-9011": (agent.BUYUK, "Büyük geldi"), "IA-9012": (agent.BUYUK, "Paçası uzun geldi"),
         "IA-9014": (agent.GEC, "Kargo çok geç geldi"), "IA-9015": (agent.KUCUK, "Dar geldi")}


def sonuc(dogru=True) -> dict:
    d = {i: {"id": i, "kategori": k, "alinti": a} for i, (k, a) in DOGRU.items()}
    if not dogru:
        d["IA-9012"]["kategori"] = agent.KALIP
        d["IA-9006"]["alinti"] = "fermuar kırık"
        d["IA-9014"]["kategori"] = agent.VAZGECTI
    return d


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    d = sonuc("<duzeltme>" in kullanici)
    return {"iadeler": [d[i] for i in re.findall(r'<iade id="([^"]+)"', kullanici)]}


class KodTesti(unittest.TestCase):
    def test_kurallar(self):
        self.assertEqual([agent.kural(x) for x in ("Bel çok dar, bir beden büyüğünü alacağım", "Paçası uzun geldi", "M sipariş ettim S beden geldi, yanlış ürün",
                                                    "Dikiş yeri sökük", "Kargo geç geldi", "Rengi farklı", "Kadar güzel değil", "Küçük geldi ama bol duruyor", "")],
                         [agent.KUCUK, agent.BUYUK, agent.YANLIS, agent.KUSUR, agent.GEC, None, None, None, None])

    def test_denetim_ve_analiz(self):
        iade = agent.iadeleri_oku(IADE)
        self.assertEqual((len(iade), sum(1 for x in iade if not x["aciklama"])), (15, 1))
        self.assertEqual(agent.denetle(iade, sonuc()), [])
        s = agent.denetle(iade, sonuc(False))
        self.assertEqual(len(s), 3)
        self.assertTrue(any("IA-9012" in x and "Beden büyük" in x for x in s) and any("fermuar kırık" in x for x in s) and any("IA-9014" in x for x in s))
        a = agent.urun_analizi(iade, sonuc(), agent.satislari_oku(SATIS), 15, 2)
        self.assertEqual([(r["urun"], r["oran"]) for r in a], [("ELB-110", 30.0), ("PNT-440", 12.0), ("MNT-220", 10.0), ("AYK-550", 2.5), ("GML-330", 1.7)])
        self.assertEqual([n.split(" (")[0].split(" >")[0] for n in a[0]["notlar"]], ["İade oranı %30", "Beden küçük", agent.RENK])
        self.assertTrue(a[1]["notlar"][0].startswith("Beden büyük (2 iade)"))
        self.assertEqual(a[1]["dagilim"][agent.BELIRSIZ], 1)
        self.assertIsNone(agent.urun_analizi(iade, sonuc(), None, 15, 2)[0]["oran"])


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "i.xlsx"
            s = agent.calistir(IADE, cikti, SATIS, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            self.assertNotIn("IA-9013", MESAJLAR[0])
            self.assertNotIn("S-50111", MESAJLAR[0])
            self.assertEqual(set(re.findall(r'<iade id="([^"]+)"', MESAJLAR[1])), {"IA-9006", "IA-9012", "IA-9014"})
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Ürün Analizi", "İadeler", "Neden Dağılımı"])
            self.assertEqual((wb["Ürün Analizi"].cell(2, 1).value, wb["Neden Dağılımı"].cell(2, 1).value), ("ELB-110", agent.KUCUK))

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(IADE, Path(tmp) / "i.xlsx", tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 3)
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--girdi", str(IADE.with_name("yok.csv")), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
