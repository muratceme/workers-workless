"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
MESAJLAR = []


def faaliyet(f, k="-"):
    return {"faaliyet": f, "sorumlu": "Proses Mühendisi", "sure": "D+5", "ilgili_kok_neden": k, "etkinlik_dogrulama": "3 lot 0 PPM"}


IYI = {
    "d1_ekip": [{"rol": "Kalite Müdürü", "gorev": "Ekip lideri"}],
    "d2_problem": {"ozet": "Kilit tırnağı kırık.", "bes_n_bir_k": {k: "x" for k in ("ne", "nerede", "ne_zaman", "kim", "neden_onemli", "nasil_ne_kadar")},
                   "is_is_not": [{"boyut": "Kalıp", "var": "K-118", "yok": "K-118B"}]},
    "d3_gecici_onlemler": [faaliyet("Müşteri stoğunu ayıkla")],
    "d4_kok_nedenler": [
        {"id": "K1", "tur": "oluşum", "kategori": "Makine", "neden": "Kalıp ayırma yüzeyinde aşınma (çapak)", "bes_neden": ["a", "b"],
         "dogrulama_durumu": "doğrulanmalı", "dogrulama_yontemi": "Kalıp ölçümü"},
        {"id": "K2", "tur": "kaçış", "kategori": "Ölçüm", "neden": "Final kontrolde tırnak testi yok", "bes_neden": ["a"],
         "dogrulama_durumu": "doğrulanmalı", "dogrulama_yontemi": "Kontrol planı incelemesi"}],
    "d5_d6_kalici_faaliyetler": [faaliyet("Kalıp revizyonu", "K1"), faaliyet("Tırnak takma testi", "K2")],
    "d7_onleme": [faaliyet("PFMEA güncelle")],
    "d8_kapanis": "Müşteri onayıyla kapanır.",
    "veri_ihtiyaci": ["Kalıp bakım kaydı"],
}


def sahte(yanit):
    def f(sistem, kullanici, sema, max_tokens=16000):
        MESAJLAR.append(kullanici)
        return copy.deepcopy(yanit)
    return f


class AnalizTesti(unittest.TestCase):
    def test_lot_ppm_ve_etkilenen(self):
        metin = (ORNEK / "musteri_sikayeti.txt").read_text(encoding="utf-8")
        a = agent.analiz_et(agent.muayene_oku(ORNEK / "muayene_kayitlari.csv"), metin)
        self.assertEqual(a["lotlar"]["L2609-14"]["hatali"], 9)            # aynı lotta iki hata türü satırı toplanır
        self.assertEqual(a["lotlar"]["L2609-14"]["ppm"], 18000)
        self.assertAlmostEqual(a["ort_ppm"], 20 / 3000 * 1e6)
        self.assertEqual(a["etkilenen"], ["L2609-14", "L2609-15"])
        self.assertEqual(a["makine"]["ENJ-04"][0], 9500)
        self.assertEqual(a["pareto"].most_common(1)[0][0], "Çapak")
        self.assertEqual(str(a["supheli_aralik"][0]), "2026-09-24")

    def test_sayi(self):
        self.assertEqual((agent.sayi("1.234"), agent.sayi("1,5"), agent.sayi("500")), (1234.0, 1.5, 500.0))


class DenetimTesti(unittest.TestCase):
    def test_iyi_taslak(self):
        self.assertEqual(agent.taslak_denetle(copy.deepcopy(IYI)), [])

    def test_eksik_kacis_ve_bagsiz_faaliyet(self):
        t = copy.deepcopy(IYI)
        t["d4_kok_nedenler"] = t["d4_kok_nedenler"][:1]
        t["d5_d6_kalici_faaliyetler"] = [faaliyet("X", "K9")]
        s = agent.taslak_denetle(t)
        self.assertIn("D4: kaçış kök nedeni yazılmamış", s)
        self.assertTrue(any("K1" in x and "bağlı kalıcı faaliyet yok" in x for x in s))
        self.assertTrue(any("K9" in x for x in s))


class AgentTesti(unittest.TestCase):
    def setUp(self):
        MESAJLAR.clear()

    def test_uctan_uca(self):
        with mock.patch.object(llm, "json_iste", side_effect=sahte(IYI)), tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "8d.xlsx"
            s = agent.calistir(ORNEK / "musteri_sikayeti.txt", cikti, ORNEK / "muayene_kayitlari.csv", ORNEK / "bilgi.json", evet=True)
            m = MESAJLAR[0]
            self.assertIn("<veri_analizi>", m)
            self.assertIn("18000 PPM", m)
            self.assertNotIn("0262 000 00 00", m)                              # telefon maskelenir
            self.assertEqual(s["sorunlar"], [])
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Özet", "D1 Ekip", "D2 Problem", "D3 Geçici Önlemler", "D4 Kök Neden",
                                             "D5-D6 Kalıcı Faaliyetler", "D7 Önleme", "Veri İhtiyacı", "Veri Analizi"])
            md = s["md"].read_text(encoding="utf-8")
            self.assertIn("**Oluşum**", md)
            self.assertIn("**Kaçış**", md)

    def test_muayenesiz(self):
        with mock.patch.object(llm, "json_iste", side_effect=sahte(IYI)), tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(ORNEK / "musteri_sikayeti.txt", Path(tmp) / "8d.xlsx", evet=True)
        self.assertEqual(s["analiz"], {})
        self.assertIn("Muayene verisi verilmedi", MESAJLAR[0])


if __name__ == "__main__":
    unittest.main()
