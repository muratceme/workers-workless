"""
8D Rapor Taslağı — Workers / Workless AI Agent
Üretim › Kalite › Kalite Mühendisi

1. Kod, muayene kayıtlarından (varsa) lot bazında hata oranını (PPM), etkilenen lotları, makine/kalıp
   kırılımını, hata türü Pareto'sunu ve şüpheli üretim tarih aralığını hesaplar (yapay zekâ yok).
2. Şikâyet metni (telefon/e-posta maskeli), parça/ekip bilgisi ve kod bulguları modele gönderilir; model D1–D8
   adımlarını içeren 8D taslağını yazar. Kök neden oluşum ve kaçış için ayrı yazılır; veriyle desteklenmeyen
   her neden "doğrulanmalı" hipotezi olarak işaretlenir ve doğrulama yöntemi verilir.
3. Kod taslağı denetler: oluşum ve kaçış kök nedeni var mı, her kök nedene D5 faaliyeti bağlanmış mı,
   doğrulanmamış neden "kesin" diye yazılmış mı. Eksikler raporda işaretlenir.
4. Çıktı: Excel (her D ayrı sayfa, 'Onay' sütunu) + Markdown taslak.

Kullanım:
    python agent.py                                              # örnek şikâyetle dener
    python agent.py --sikayet sikayet.txt --muayene muayene.xlsx --bilgi bilgi.json
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import Counter, OrderedDict, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
ALTI_M = ["İnsan", "Makine", "Metot", "Malzeme", "Ölçüm", "Çevre"]


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def sayi(x) -> float:
    if x in (None, ""):
        return 0.0
    if isinstance(x, (int, float)):
        return float(x)
    t = str(x).strip().replace(" ", "")
    if "," in t:
        t = t.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(\.\d{3})+", t):         # 1.234 → binlik ayırıcı
        t = t.replace(".", "")
    try:
        return float(t)
    except ValueError:
        return 0.0


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(str(x).strip(), f).date()
        except ValueError:
            pass
    return None


# ----------------------------------------------------------------------------
# Veri analizi (deterministik)
# ----------------------------------------------------------------------------

def muayene_oku(yol: Path) -> list[dict]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    s = [r for r in s if any(c not in (None, "") for c in r)]
    b = [kucuk(x) for x in s[0]]
    bul = lambda *a: next((i for i, x in enumerate(b) if x in a), None)  # noqa: E731
    k = {"tarih": bul("tarih", "üretim tarihi"), "lot": bul("lot", "parti", "lot no"), "makine": bul("makine", "hat", "tezgâh", "tezgah"),
         "kalip": bul("kalıp", "kalip", "aparat"), "kontrol": bul("kontrol edilen", "muayene edilen", "örnek"),
         "hatali": bul("hatalı", "hatali", "red", "ret"), "tur": bul("hata türü", "hata", "kusur")}
    if k["lot"] is None or k["hatali"] is None:
        raise SystemExit(f"Muayene kayıtlarında Lot ve Hatalı sütunları gerekli. Başlıklar: {s[0]}")
    al = lambda r, a: r[k[a]] if k[a] is not None and k[a] < len(r) else None  # noqa: E731
    return [{"tarih": tarih(al(r, "tarih")), "lot": str(al(r, "lot") or "").strip(), "makine": str(al(r, "makine") or "").strip(),
             "kalip": str(al(r, "kalip") or "").strip(), "kontrol": sayi(al(r, "kontrol")), "hatali": sayi(al(r, "hatali")),
             "tur": str(al(r, "tur") or "").strip()} for r in s[1:]]


def analiz_et(kayitlar: list[dict], sikayet_metni: str) -> dict:
    lotlar: "OrderedDict[str, dict]" = OrderedDict()
    for r in sorted(kayitlar, key=lambda r: (r["tarih"] or date.min, r["lot"])):
        x = lotlar.setdefault(r["lot"], {"lot": r["lot"], "tarihler": set(), "makine": set(), "kalip": set(), "kontrol": 0.0, "hatali": 0.0,
                                         "turler": Counter()})
        if r["tarih"]:
            x["tarihler"].add(r["tarih"])
        if r["makine"]:
            x["makine"].add(r["makine"])
        if r["kalip"]:
            x["kalip"].add(r["kalip"])
        x["kontrol"] += r["kontrol"]
        x["hatali"] += r["hatali"]
        if r["hatali"] and r["tur"]:
            x["turler"][r["tur"]] += r["hatali"]
    for x in lotlar.values():
        x["ppm"] = x["hatali"] / x["kontrol"] * 1e6 if x["kontrol"] else None
    toplam_k = sum(x["kontrol"] for x in lotlar.values())
    toplam_h = sum(x["hatali"] for x in lotlar.values())
    ort_ppm = toplam_h / toplam_k * 1e6 if toplam_k else None
    sikayet_lotlari = [l for l in lotlar if l and re.search(r"(?<![\w-])" + re.escape(l) + r"(?![\w-])", sikayet_metni)]
    # Etkilenen: şikâyette geçen lotlar + ortalamanın 2 katından yüksek PPM'li lotlar
    yuksek = [l for l, x in lotlar.items() if x["ppm"] is not None and ort_ppm and x["ppm"] > 2 * ort_ppm]
    etkilenen = list(dict.fromkeys(sikayet_lotlari + yuksek))
    tarihler = sorted(t for l in etkilenen for t in lotlar[l]["tarihler"])

    def kirilim(alan):
        k = defaultdict(lambda: [0.0, 0.0])
        for x in lotlar.values():
            for v in x[alan] or {"(belirtilmemiş)"}:
                k[v][0] += x["kontrol"] / max(1, len(x[alan]))
                k[v][1] += x["hatali"] / max(1, len(x[alan]))
        return {v: (h / kk * 1e6 if kk else None, kk, h) for v, (kk, h) in k.items()}

    pareto = Counter()
    for x in lotlar.values():
        pareto.update(x["turler"])
    return {"lotlar": lotlar, "ort_ppm": ort_ppm, "toplam_kontrol": toplam_k, "toplam_hatali": toplam_h,
            "sikayet_lotlari": sikayet_lotlari, "etkilenen": etkilenen,
            "supheli_aralik": (tarihler[0], tarihler[-1]) if tarihler else None,
            "makine": kirilim("makine"), "kalip": kirilim("kalip"), "pareto": pareto}


def analiz_metni(a: dict) -> str:
    if not a:
        return "Muayene verisi verilmedi."
    s = [f"- Toplam: {a['toplam_kontrol']:.0f} adet kontrol, {a['toplam_hatali']:.0f} hatalı, ortalama {a['ort_ppm']:.0f} PPM"]
    for l, x in a["lotlar"].items():
        s.append(f"- Lot {l} ({', '.join(t.strftime('%d.%m.%Y') for t in sorted(x['tarihler']))}; makine {', '.join(sorted(x['makine'])) or '-'}; "
                 f"kalıp {', '.join(sorted(x['kalip'])) or '-'}): {x['kontrol']:.0f} kontrol, {x['hatali']:.0f} hatalı, "
                 f"{x['ppm']:.0f} PPM" if x["ppm"] is not None else f"- Lot {l}: kontrol adedi yok")
    s.append(f"- Şikâyette geçen lotlar: {', '.join(a['sikayet_lotlari']) or '-'}")
    s.append(f"- Etkilenen lotlar (şikâyette geçen + ortalamanın 2 katından yüksek PPM): {', '.join(a['etkilenen']) or '-'}")
    if a["supheli_aralik"]:
        s.append(f"- Şüpheli üretim aralığı: {a['supheli_aralik'][0]:%d.%m.%Y} – {a['supheli_aralik'][1]:%d.%m.%Y}")
    for ad in ("makine", "kalip"):
        s.append(f"- {ad.capitalize()} kırılımı: " + "; ".join(f"{k}: {v[0]:.0f} PPM ({v[2]:.0f}/{v[1]:.0f})" for k, v in a[ad].items()
                                                             if v[0] is not None))
    s.append("- Hata türü Pareto: " + ", ".join(f"{t} {n:.0f}" for t, n in a["pareto"].most_common()))
    return "\n".join(s)


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

_KOK = {
    "type": "object",
    "properties": {
        "id": {"type": "string"},
        "tur": {"type": "string", "enum": ["oluşum", "kaçış"]},
        "kategori": {"type": "string", "enum": ALTI_M},
        "neden": {"type": "string"},
        "bes_neden": {"type": "array", "items": {"type": "string"}},
        "dogrulama_durumu": {"type": "string", "enum": ["doğrulandı", "doğrulanmalı"]},
        "dogrulama_yontemi": {"type": "string"},
    },
    "required": ["id", "tur", "kategori", "neden", "bes_neden", "dogrulama_durumu", "dogrulama_yontemi"],
    "additionalProperties": False,
}
_FAALIYET = {
    "type": "object",
    "properties": {"faaliyet": {"type": "string"}, "sorumlu": {"type": "string"}, "sure": {"type": "string"},
                   "ilgili_kok_neden": {"type": "string"}, "etkinlik_dogrulama": {"type": "string"}},
    "required": ["faaliyet", "sorumlu", "sure", "ilgili_kok_neden", "etkinlik_dogrulama"],
    "additionalProperties": False,
}
SEMA = {
    "type": "object",
    "properties": {
        "d1_ekip": {"type": "array", "items": {"type": "object", "properties": {"rol": {"type": "string"}, "gorev": {"type": "string"}},
                                                "required": ["rol", "gorev"], "additionalProperties": False}},
        "d2_problem": {"type": "object", "properties": {
            "ozet": {"type": "string"},
            "bes_n_bir_k": {"type": "object", "properties": {k: {"type": "string"} for k in ("ne", "nerede", "ne_zaman", "kim", "neden_onemli", "nasil_ne_kadar")},
                            "required": ["ne", "nerede", "ne_zaman", "kim", "neden_onemli", "nasil_ne_kadar"], "additionalProperties": False},
            "is_is_not": {"type": "array", "items": {"type": "object", "properties": {"boyut": {"type": "string"}, "var": {"type": "string"}, "yok": {"type": "string"}},
                                                     "required": ["boyut", "var", "yok"], "additionalProperties": False}}},
            "required": ["ozet", "bes_n_bir_k", "is_is_not"], "additionalProperties": False},
        "d3_gecici_onlemler": {"type": "array", "items": _FAALIYET},
        "d4_kok_nedenler": {"type": "array", "items": _KOK},
        "d5_d6_kalici_faaliyetler": {"type": "array", "items": _FAALIYET},
        "d7_onleme": {"type": "array", "items": _FAALIYET},
        "d8_kapanis": {"type": "string"},
        "veri_ihtiyaci": {"type": "array", "items": {"type": "string"}},
    },
    "required": ["d1_ekip", "d2_problem", "d3_gecici_onlemler", "d4_kok_nedenler", "d5_d6_kalici_faaliyetler", "d7_onleme", "d8_kapanis", "veri_ihtiyaci"],
    "additionalProperties": False,
}


def taslak_denetle(t: dict) -> list[str]:
    sorunlar = []
    turler = {k["tur"] for k in t["d4_kok_nedenler"]}
    for tur in ("oluşum", "kaçış"):
        if tur not in turler:
            sorunlar.append(f"D4: {tur} kök nedeni yazılmamış")
    idler = {k["id"] for k in t["d4_kok_nedenler"]}
    bagli = {i.strip() for f in t["d5_d6_kalici_faaliyetler"] for i in re.split(r"[,;\s]+", f["ilgili_kok_neden"]) if i.strip()}
    for k in t["d4_kok_nedenler"]:
        if k["id"] not in bagli:
            sorunlar.append(f"D5: {k['id']} ('{k['neden'][:50]}') kök nedenine bağlı kalıcı faaliyet yok")
    for f in t["d5_d6_kalici_faaliyetler"]:
        bilinmeyen = {i.strip() for i in re.split(r"[,;\s]+", f["ilgili_kok_neden"]) if i.strip()} - idler
        if bilinmeyen:
            sorunlar.append(f"D5: '{f['faaliyet'][:50]}' faaliyeti olmayan kök nedene bağlanmış ({', '.join(sorted(bilinmeyen))})")
    for k in t["d4_kok_nedenler"]:
        if k["dogrulama_durumu"] == "doğrulandı" and not k["dogrulama_yontemi"].strip():
            sorunlar.append(f"D4: '{k['neden'][:60]}' doğrulandı denmiş ama doğrulama kanıtı/yöntemi yok")
    if not t["d3_gecici_onlemler"]:
        sorunlar.append("D3: geçici önlem yok")
    return sorunlar


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
ONAY_DOLGU = PatternFill("solid", fgColor="FFF4CE")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
UST = Alignment(vertical="top", wrap_text=True)


def _sayfa(wb, ad, basliklar, satirlar, genislik, onay=True):
    ws = wb.create_sheet(ad)
    ws.append(basliklar + (["Onay"] if onay else []))
    for h in ws[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for s in satirlar:
        ws.append(list(s) + ([""] if onay else []))
        for h in ws[ws.max_row]:
            h.alignment = UST
        if onay:
            ws.cell(ws.max_row, len(basliklar) + 1).fill = ONAY_DOLGU
    for j, w in enumerate(list(genislik) + ([10] if onay else []), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    return ws


def rapor_yaz(cikti: Path, bilgi: dict, analiz: dict, t: dict, sorunlar: list[str]) -> Path:
    wb = Workbook()
    o = wb.active
    o.title = "Özet"
    p = t["d2_problem"]
    for s in [["8D RAPORU — TASLAK"], [], ["Müşteri", bilgi.get("musteri", "")], ["Parça", f"{bilgi.get('parca_no', '')} {bilgi.get('parca_adi', '')}".strip()],
              ["Şikâyet no", bilgi.get("sikayet_no", "")], ["Tespit tarihi", bilgi.get("tespit_tarihi", "")],
              ["Hedef süreler", ", ".join(f"{k}: {v}" for k, v in bilgi.get("hedef_sureler", {}).items())], [],
              ["Problem (D2)", p["ozet"]], ["D8 kapanış", t["d8_kapanis"]],
              ["Taslak kontrolü", "; ".join(sorunlar) or "Oluşum ve kaçış kök nedenleri ve bağlı faaliyetler mevcut"],
              ["Model", llm.kullanim_ozeti()], ["Oluşturulma", datetime.now().strftime("%d.%m.%Y %H:%M")],
              ["Önemli", "Yapay zekâ taslağıdır. 'doğrulanmalı' işaretli kök nedenler hipotezdir; veri toplanıp doğrulanmadan "
                         "müşteriye kesin kök neden olarak bildirilmemelidir."]]:
        o.append(s)
    o["A1"].font = Font(bold=True, size=13)
    if sorunlar:
        o.cell(11, 2).fill = KIRMIZI
    o.column_dimensions["A"].width = 18
    o.column_dimensions["B"].width = 120
    for satir in o.iter_rows():
        for h in satir:
            h.alignment = UST

    _sayfa(wb, "D1 Ekip", ["Rol", "Görev"], [(e["rol"], e["gorev"]) for e in t["d1_ekip"]], (36, 80))
    d2 = [(k.replace("_", " ").capitalize(), v) for k, v in p["bes_n_bir_k"].items()]
    ws = _sayfa(wb, "D2 Problem", ["5N1K", "Açıklama"], d2, (20, 110), onay=False)
    ws.append([])
    ws.append(["Is / Is Not", "VAR (problem görülen)", "YOK (görülmeyen)"])
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for x in p["is_is_not"]:
        ws.append([x["boyut"], x["var"], x["yok"]])
        for h in ws[ws.max_row]:
            h.alignment = UST
    ws.column_dimensions["C"].width = 60
    fa = lambda liste: [(f["faaliyet"], f["sorumlu"], f["sure"], f["ilgili_kok_neden"], f["etkinlik_dogrulama"]) for f in liste]  # noqa: E731
    fb = ["Faaliyet", "Sorumlu (rol)", "Süre", "İlgili Kök Neden", "Etkinlik Doğrulama"]
    _sayfa(wb, "D3 Geçici Önlemler", fb, fa(t["d3_gecici_onlemler"]), (60, 26, 10, 36, 44))
    k = _sayfa(wb, "D4 Kök Neden", ["No", "Tür", "6M", "Neden", "5 Neden Zinciri", "Doğrulama Durumu", "Doğrulama Yöntemi"],
               [(x["id"], x["tur"], x["kategori"], x["neden"], " → ".join(x["bes_neden"]), x["dogrulama_durumu"], x["dogrulama_yontemi"])
                for x in sorted(t["d4_kok_nedenler"], key=lambda x: (x["tur"], ALTI_M.index(x["kategori"])))],
               (6, 9, 10, 40, 70, 14, 44))
    for r in range(2, k.max_row + 1):
        if k.cell(r, 6).value == "doğrulanmalı":
            k.cell(r, 6).fill = ONAY_DOLGU
    _sayfa(wb, "D5-D6 Kalıcı Faaliyetler", fb, fa(t["d5_d6_kalici_faaliyetler"]), (60, 26, 10, 36, 44))
    _sayfa(wb, "D7 Önleme", fb, fa(t["d7_onleme"]), (60, 26, 10, 36, 44))
    _sayfa(wb, "Veri İhtiyacı", ["Toplanacak veri / kanıt"], [(v,) for v in t["veri_ihtiyaci"]], (110,))

    if analiz:
        v = wb.create_sheet("Veri Analizi")
        v.append(["Lot", "Tarih", "Makine", "Kalıp", "Kontrol", "Hatalı", "PPM", "Hata Türleri", "Etkilenen"])
        for h in v[1]:
            h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        for l, x in analiz["lotlar"].items():
            v.append([l, ", ".join(t_.strftime("%d.%m.%Y") for t_ in sorted(x["tarihler"])), ", ".join(sorted(x["makine"])),
                      ", ".join(sorted(x["kalip"])), x["kontrol"], x["hatali"], None if x["ppm"] is None else round(x["ppm"]),
                      ", ".join(f"{a} {n:.0f}" for a, n in x["turler"].most_common()), "EVET" if l in analiz["etkilenen"] else ""])
            if l in analiz["etkilenen"]:
                v.cell(v.max_row, 9).fill = KIRMIZI
        for j, w in enumerate((12, 14, 10, 10, 9, 8, 9, 34, 10), 1):
            v.column_dimensions[get_column_letter(j)].width = w
        n = len(analiz["lotlar"])
        if n:
            g = BarChart()
            g.title, g.height, g.width = "Lot bazında PPM", 8, 18
            g.add_data(Reference(v, min_col=7, min_row=1, max_row=1 + n), titles_from_data=True)
            g.set_categories(Reference(v, min_col=1, min_row=2, max_row=1 + n))
            v.add_chart(g, "K2")
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

    md = [f"# 8D raporu (taslak) — {bilgi.get('parca_no', '')} {bilgi.get('parca_adi', '')}".rstrip(), "",
          f"**Müşteri:** {bilgi.get('musteri', '-')} · **Şikâyet no:** {bilgi.get('sikayet_no', '-')} · **Tespit:** {bilgi.get('tespit_tarihi', '-')}", "",
          "## D1 Ekip", *[f"- {e['rol']}: {e['gorev']}" for e in t["d1_ekip"]], "", "## D2 Problem tanımı", p["ozet"], "",
          *[f"- **{a}:** {b}" for a, b in d2], "", "| Boyut | VAR | YOK |", "|---|---|---|",
          *[f"| {x['boyut']} | {x['var']} | {x['yok']} |" for x in p["is_is_not"]], "", "## D3 Geçici önlemler",
          *[f"- {f['faaliyet']} — _{f['sorumlu']}, {f['sure']}_" for f in t["d3_gecici_onlemler"]], "", "## D4 Kök neden"]
    for tur in ("oluşum", "kaçış"):
        md.append(f"**{tur.capitalize()}**")
        md += [f"- {x['id']} [{x['kategori']}] {x['neden']} — _{x['dogrulama_durumu']}_ ({x['dogrulama_yontemi']})"
               for x in t["d4_kok_nedenler"] if x["tur"] == tur]
    md += ["", "## D5–D6 Kalıcı düzeltici faaliyetler", *[f"- {f['faaliyet']} — _{f['sorumlu']}, {f['sure']}_; doğrulama: {f['etkinlik_dogrulama']}"
                                                           for f in t["d5_d6_kalici_faaliyetler"]],
           "", "## D7 Önleme", *[f"- {f['faaliyet']} — _{f['sorumlu']}_" for f in t["d7_onleme"]], "", "## D8 Kapanış", t["d8_kapanis"], "",
           "## Toplanacak veri", *[f"- {v_}" for v_ in t["veri_ihtiyaci"]], "", "---",
           "_Yapay zekâ taslağı: veri analizi koddan, 8D içeriği modelden. 'doğrulanmalı' kök nedenler hipotezdir._"]
    if sorunlar:
        md.insert(3, "> **Taslak kontrolü:** " + "; ".join(sorunlar) + "\n")
    md_yol = cikti.with_suffix(".md")
    md_yol.write_text("\n".join(md), encoding="utf-8")
    return md_yol


def calistir(sikayet: Path, cikti: Path, muayene: Path | None = None, bilgi_yolu: Path | None = None, evet: bool = False) -> dict:
    metin = belge.metin_oku(sikayet)
    bilgi = json.loads(bilgi_yolu.read_text(encoding="utf-8")) if bilgi_yolu else {}
    analiz = analiz_et(muayene_oku(muayene), metin) if muayene else {}
    if analiz:
        print(f"[OK] Veri analizi: {len(analiz['lotlar'])} lot · ortalama {analiz['ort_ppm']:.0f} PPM · etkilenen: {', '.join(analiz['etkilenen']) or '-'}")
    llm.onay_al("Şikâyet metni (telefon/e-posta maskeli), parça/ekip bilgisi ve muayene veri özeti 8D taslağı için gönderilecek.", evet)
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    mesaj = "\n".join([f"<sikayet>\n{llm.maskele(metin)}\n</sikayet>",
                       f"<bilgi>\n{json.dumps(bilgi, ensure_ascii=False, indent=1)}\n</bilgi>",
                       f"<veri_analizi>\n{analiz_metni(analiz)}\n</veri_analizi>"])
    taslak = llm.json_iste(sistem, mesaj, SEMA)
    sorunlar = taslak_denetle(taslak)
    md = rapor_yaz(cikti, bilgi, analiz, taslak, sorunlar)
    return {"analiz": analiz, "taslak": taslak, "sorunlar": sorunlar, "md": md}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Müşteri şikâyeti / uygunsuzluk kaydından 8D raporu taslağı hazırlar.")
    p.add_argument("--sikayet", type=Path, default=BURASI / "ornek_veri" / "musteri_sikayeti.txt", help="Şikâyet / uygunsuzluk metni (.txt, .docx, .pdf)")
    p.add_argument("--muayene", type=Path, help="Muayene kayıtları (.xlsx/.csv: Tarih, Lot, Makine, Kalıp, Kontrol Edilen, Hatalı, Hata Türü)")
    p.add_argument("--bilgi", type=Path, help="Parça, müşteri, ekip, hedef süreler (JSON)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "8d_taslak.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    if a.sikayet == BURASI / "ornek_veri" / "musteri_sikayeti.txt":
        a.muayene = a.muayene or BURASI / "ornek_veri" / "muayene_kayitlari.csv"
        a.bilgi = a.bilgi or BURASI / "ornek_veri" / "bilgi.json"
    try:
        s = calistir(a.sikayet, a.cikti, a.muayene, a.bilgi, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] 8D taslağı: {len(s['taslak']['d4_kok_nedenler'])} kök neden, {len(s['taslak']['d5_d6_kalici_faaliyetler'])} kalıcı faaliyet")
    for x in s["sorunlar"]:
        print(f"[!] {x}")
    print(f"[OK] Rapor: {a.cikti.resolve()} (+ {s['md'].name})")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
