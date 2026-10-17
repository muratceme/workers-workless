"""
Serbest Tariften Standart Reçete Kartı — Gerisi Bende AI Agent
Otel › Yiyecek-İçecek › Aşçıbaşı

Şeflerin serbest yazdığı tarifleri standart reçete kartına dönüştürür: malzeme, miktar, ölçü, adımlar, alerjenler.
1. Kod tarifleri ("Tarif: Ad | Porsiyon: N" başlıklı bloklar), fiyat listesini ("Porsiyon Maliyeti Hesaplama" kod
   bloğunun dosyası) ve ölçü dönüşüm tablosunu okur.
2. Model malzemeleri tarifte yazdığı ölçüyle (su bardağı, yemek kaşığı, adet …), adımları ve alerjenleri çıkarır;
   malzeme adını fiyat listesindeki adla eşleştirir.
3. Kod denetler: tarifteki her malzeme satırı alınmış mı; miktar ve ölçü tarifte yazıyor mu; kanıt tarifte aynen var mı;
   eşleştirilen ad tarifteki adla uyumlu mu (tereyağı ≠ ayçiçek yağı); adımlarda tarifte olmayan sayı var mı; kodun
   anahtar kelimelerle bulduğu alerjenler modelin listesinde mi. Sorun varsa yalnız o tarif modele geri gönderilir.
4. Kod ev ölçülerini dönüşüm tablosuyla g / ml'ye çevirir, porsiyon başı miktarı hesaplar ve kod bloğunun reçete ve
   menü dosyalarını yazar.

Kullanım:
    python agent.py                                                            # örnek 2 tarif
    python agent.py --tarifler tarifler.docx --fiyatlar fiyatlar.xlsx --donusum olcu_donusum.xlsx
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KOD_DOLGU = PatternFill("solid", fgColor="E8F0FE")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}
OLCULER = ["g", "kg", "ml", "l", "adet", "su bardağı", "çay bardağı", "yemek kaşığı", "tatlı kaşığı", "çay kaşığı", "tutam", "diş", "demet"]
DOGRUDAN = {"g": ("g", 1), "kg": ("g", 1000), "ml": ("ml", 1), "l": ("ml", 1000)}
OLCU_YAZIM = {"g": ("g", "gr", "gram"), "kg": ("kg", "kilo"), "ml": ("ml",), "l": ("l", "lt", "litre")}
SOZ = {"yarim": Decimal("0.5"), "ceyrek": Decimal("0.25"), "bir": Decimal(1), "iki": Decimal(2), "uc": Decimal(3), "dort": Decimal(4), "bes": Decimal(5)}
# 14 alerjen grubu (AB 1169/2011 Ek II; Türk Gıda Kodeksi etiketleme yönetmeliğinde de aynı gruplar) ve kodun aradığı kelimeler
ALERJEN = {
    "Gluten içeren tahıllar": ("un", "bugday", "bulgur", "irmik", "ekmek", "galeta", "makarna", "arpa", "cavdar", "yulaf", "sehriye", "yufka", "kadayif"),
    "Kabuklular": ("karides", "yengec", "istakoz", "kerevit"),
    "Yumurta": ("yumurta",),
    "Balık": ("balik", "hamsi", "levrek", "cupra", "somon", "ton baligi", "palamut", "sardalya"),
    "Yer fıstığı": ("yer fistigi",),
    "Soya": ("soya",),
    "Süt ve süt ürünleri (laktoz dahil)": ("sut", "tereyag", "yogurt", "peynir", "krema", "kaymak", "labne", "ayran", "kefir"),
    "Sert kabuklu meyveler": ("findik", "ceviz", "badem", "antep fistigi", "kaju", "pikan", "macadamia"),
    "Kereviz": ("kereviz",),
    "Hardal": ("hardal",),
    "Susam": ("susam", "tahin"),
    "Kükürt dioksit ve sülfitler": ("sulfit", "sarap"),
    "Acı bakla (lupin)": ("aci bakla", "lupin"),
    "Yumuşakçalar": ("midye", "kalamar", "ahtapot", "supya", "istiridye"),
}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def sayi(x) -> Decimal | None:
    s = re.sub(r"[^\d,.\-/]", "", str(x if x is not None else ""))
    if re.fullmatch(r"\d+/\d+", s):
        a, b = s.split("/")
        return Decimal(a) / Decimal(b) if int(b) else None
    s = s.replace(",", ".")
    try:
        return Decimal(s) if s not in ("", "-", ".") else None
    except InvalidOperation:
        return None


def metin_sayilari(metin: str) -> set[Decimal]:
    out = {v for x in re.findall(r"\d+(?:/\d+|[.,]\d+)?", str(metin)) if (v := sayi(x)) is not None}
    k = katla(metin)
    out |= {v for w, v in SOZ.items() if re.search(rf"\b{w}\b", k)}
    return out


def kok(s: str) -> set[str]:
    return {w[:4] for w in katla(s).split() if len(w) >= 2 and not w.isdigit()}


def alerjen_bul(metin: str) -> set[str]:
    k = katla(metin)
    return {a for a, ks in ALERJEN.items() if any(re.search(rf"\b{re.escape(x)}", k) for x in ks)}


# ----------------------------------------------------------------------------
# Girdi
# ----------------------------------------------------------------------------

def tablo_oku(yol: Path) -> list[dict]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, read_only=True, data_only=True)
        s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=metin.split("\n")[0].count)))
    s = [r for r in s if any(c not in (None, "") for c in r)]
    bas = [katla(c) for c in s[0]] if s else []
    return [dict(zip(bas, list(r) + [None] * (len(bas) - len(r)))) for r in s[1:]]


def fiyat_listesi(yol: Path | None) -> dict[str, str]:
    """katlanmış ad → fiyat listesindeki ad, birim"""
    if not yol:
        return {}
    return {katla(r.get("malzeme")): (str(r.get("malzeme")).strip(), katla(r.get("birim"))) for r in tablo_oku(yol) if r.get("malzeme")}


def donusum_oku(yol: Path | None) -> dict[tuple[str, str], tuple[Decimal, str]]:
    out = {}
    if yol:
        for r in tablo_oku(yol):
            o, mz, mk, b = katla(r.get("olcu")), katla(r.get("malzeme")) or "*", sayi(r.get("miktar")), katla(r.get("birim"))
            if o and mk and b in ("g", "ml"):
                out[(o, "*" if mz in ("", "*") else mz)] = (mk, b)
    return out


def tarifleri_oku(yol: Path) -> list[dict]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: tarif dosyası bulunamadı.")
    out, son = [], None
    for s in belge.metin_oku(yol).replace("\r", "").splitlines():
        m = re.match(r"^\s*Tarif\s*:\s*(.+)$", s, re.I)
        if m:
            alan = [x.strip() for x in m[1].split("|")]
            ek = {katla(a.split(":", 1)[0]): a.split(":", 1)[1].strip() for a in alan[1:] if ":" in a}
            son = {"ad": alan[0], "porsiyon": sayi(ek.get("porsiyon")), "satirlar": [], "hazirlik": "", "bolum": "malzeme"}
            out.append(son)
        elif son and s.strip():
            h = re.match(r"^\s*(Hazırlanışı|Hazırlanış|Yapılışı|Hazırlık)\s*:?\s*(.*)$", s, re.I)
            if h:
                son["bolum"] = "hazirlik"
                son["hazirlik"] = h[2].strip()
            elif son["bolum"] == "hazirlik":
                son["hazirlik"] = (son["hazirlik"] + " " + s.strip()).strip()
            else:
                son["satirlar"].append(s.strip())
    out = [t for t in out if t["satirlar"]]
    if not out:
        raise llm.LLMHatasi(f"{yol.name}: 'Tarif: Ad | Porsiyon: N' başlıklı tarif yok.")
    for t in out:
        t["metin"] = llm.maskele("\n".join(t["satirlar"]) + ("\nHazırlanışı: " + t["hazirlik"] if t["hazirlik"] else ""))
        del t["bolum"]
    return out


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

SEMA = {"type": "object", "properties": {
    "malzemeler": {"type": "array", "items": {"type": "object", "properties": {
        "tarifteki_ad": S, "malzeme": S, "miktar": S, "olcu": {"type": "string", "enum": OLCULER}, "fire": S, "kanit": S},
        "required": ["tarifteki_ad", "malzeme", "miktar", "olcu", "fire", "kanit"], "additionalProperties": False}},
    "adimlar": {"type": "array", "items": S}, "alerjenler": {"type": "array", "items": {"type": "string", "enum": list(ALERJEN)}},
    "belirsizlikler": {"type": "array", "items": S}},
    "required": ["malzemeler", "adimlar", "alerjenler", "belirsizlikler"], "additionalProperties": False}


def denetle(t: dict, c: dict, fiyat: dict) -> list[str]:
    ad, km, sorun = t["ad"], katla(t["metin"]), []
    ns = metin_sayilari(t["metin"])
    kanitlar = []
    for i, x in enumerate(c["malzemeler"], 1):
        on = f"{ad}: malzeme {i} ('{x['tarifteki_ad']}')"
        kk = katla(x["kanit"])
        if len(kk.split()) < 1 or kk not in km:
            sorun.append(f"{on}: kanıt tarifte aynen geçmiyor")
            continue
        kanitlar.append(kk)
        if x["miktar"].strip():
            v = sayi(x["miktar"])
            if v is None or v not in metin_sayilari(x["kanit"]):
                sorun.append(f"{on}: miktar '{x['miktar']}' kanıtta yazmıyor")
        if x["olcu"] != "adet":
            yazim = OLCU_YAZIM.get(x["olcu"], (katla(x["olcu"]),))
            if not any(re.search(rf"(?<![a-z0-9]){re.escape(katla(y))}(?![a-z])", kk) or re.search(rf"\d{re.escape(katla(y))}\b", kk) for y in yazim):
                sorun.append(f"{on}: ölçü '{x['olcu']}' kanıtta yazmıyor; tarifte yazan ölçüyü aynen kullanın")
        if not kok(x["tarifteki_ad"]) & kok(x["malzeme"]):
            sorun.append(f"{on}: '{x['malzeme']}' tarifteki adla uyuşmuyor; farklı ürünle eşleştirmeyin")
        if x["fire"].strip() and (sayi(x["fire"]) is None or sayi(x["fire"]) not in ns):
            sorun.append(f"{on}: fire '{x['fire']}' tarifte yazmıyor")
    for s in t["satirlar"]:
        ks = katla(s)
        if not any(x in ks or ks in x for x in kanitlar):
            sorun.append(f"{ad}: tarif satırı malzemelere alınmamış: \"{s[:80]}\"")
    fazla = sorted(set().union(*(metin_sayilari(a) for a in c["adimlar"])) - ns, key=float) if c["adimlar"] else []
    if fazla:
        sorun.append(f"{ad}: adımlarda tarifte olmayan sayı (süre / sıcaklık): {', '.join(f'{x:g}' for x in fazla)}")
    if t["hazirlik"] and not c["adimlar"]:
        sorun.append(f"{ad}: hazırlanış adımları boş")
    eksik = sorted(alerjen_bul(" ".join(x["tarifteki_ad"] + " " + x["malzeme"] for x in c["malzemeler"]) + " " + t["metin"]) - set(c["alerjenler"]))
    if eksik:
        sorun.append(f"{ad}: alerjen eksik: {', '.join(eksik)}")
    return sorun


def uret(tarifler: list[dict], fiyat: dict, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    ml = ", ".join(v[0] for v in fiyat.values()) or "-"
    sonuc = dict(onceki or {})
    hedef = tarifler if duzeltme is None else [t for t in tarifler if any(s.startswith(t["ad"] + ":") for s in duzeltme)]
    for t in hedef:
        p = f"{t['porsiyon']:g}" if t["porsiyon"] else "-"
        mesaj = f"<malzeme_listesi>{ml}</malzeme_listesi>\n<tarif ad=\"{t['ad']}\" porsiyon=\"{p}\">\n{t['metin']}\n</tarif>"
        d = [s.split(": ", 1)[1] for s in (duzeltme or []) if s.startswith(t["ad"] + ":")]
        if d:
            mesaj += ("\n<duzeltme>\nÖnceki cevap: " + json.dumps(sonuc.get(t["ad"]), ensure_ascii=False) + "\nKodun bulduğu sorunlar:\n"
                      + "\n".join(f"- {s}" for s in d) + "\n</duzeltme>")
        sonuc[t["ad"]] = llm.json_iste(sistem, mesaj, SEMA)
    return sonuc


def duzeltme_dongusu(uret_f, denetle_f, tur: int) -> tuple[dict, list[str], int]:
    taslak = uret_f(None, None)
    kalan = denetle_f(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret_f(kalan, taslak)
        kalan = denetle_f(taslak)
    return taslak, kalan, yapilan


# ----------------------------------------------------------------------------
# Hesap ve çıktı
# ----------------------------------------------------------------------------

def cevir(x: dict, donusum: dict, fiyat: dict) -> dict:
    """Ev ölçüsünü g / ml'ye çevirir (kod). Fiyat listesi 'adet' ise adet olarak bırakır."""
    mk = sayi(x["miktar"])
    fb = fiyat.get(katla(x["malzeme"]), (None, ""))[1]
    if mk is None:
        return {"miktar": None, "birim": "", "not": "Miktar yok; tartıp reçeteye yazın"}
    if x["olcu"] in DOGRUDAN:
        b, k = DOGRUDAN[x["olcu"]]
        return {"miktar": mk * k, "birim": b, "not": ""}
    if x["olcu"] == "adet" and fb == "adet":
        return {"miktar": mk, "birim": "adet", "not": ""}
    o = katla(x["olcu"])
    d = donusum.get((o, katla(x["malzeme"]))) or donusum.get((o, "*"))
    if d:
        return {"miktar": mk * d[0], "birim": d[1], "not": f"{mk:g} {x['olcu']} × {d[0]:g} {d[1]} (dönüşüm tablosu, yaklaşık)"}
    return {"miktar": None, "birim": "", "not": f"'{x['olcu']}' için dönüşüm yok; tablosuna ekleyin ya da tartın"}


def tablo(ws, bas: list[str], satirlar: list[list], ai: set[int] = frozenset(), kod: set[int] = frozenset(), genis: dict | None = None) -> None:
    ws.append(bas)
    for h in ws[1]:
        h.fill, h.font, h.alignment = BASLIK_DOLGU, BASLIK_YAZI, UST
    for r in satirlar:
        ws.append(r)
        for j, v in enumerate(r, 1):
            c = ws.cell(ws.max_row, j)
            c.alignment = UST
            if j in ai and v not in (None, ""):
                c.fill = AI_DOLGU
            if j in kod and v not in (None, ""):
                c.fill = KOD_DOLGU
    for j in range(1, len(bas) + 1):
        ws.column_dimensions[get_column_letter(j)].width = (genis or {}).get(j, 14)
    ws.freeze_panes = "A2"


def f(x: Decimal | None) -> float | None:
    return None if x is None else float(round(x, 2))


def rapor_yaz(cikti: Path, tarifler: list[dict], sonuc: dict, cevrim: dict, fiyat: dict, kalan: list[str], tur: int) -> tuple[Path, Path]:
    wb = Workbook()
    ws = wb.active
    ws.title = "Reçete Kartları"
    rr, blok = [], []
    for t in tarifler:
        for x, cv in zip(sonuc[t["ad"]]["malzemeler"], cevrim[t["ad"]]):
            p = cv["miktar"] / t["porsiyon"] if cv["miktar"] is not None and t["porsiyon"] else None
            listede = "Evet" if katla(x["malzeme"]) in fiyat else "Hayır"
            rr.append([t["ad"], t["porsiyon"] and int(t["porsiyon"]), x["tarifteki_ad"], x["malzeme"], x["miktar"], x["olcu"], f(cv["miktar"]), cv["birim"],
                       f(p), x["fire"], listede, cv["not"]])
            if cv["miktar"] is not None:
                blok.append([t["ad"], fiyat.get(katla(x["malzeme"]), (x["malzeme"],))[0], f(cv["miktar"]), cv["birim"], sayi(x["fire"]) and float(sayi(x["fire"]))])
    tablo(ws, ["Menü Kalemi", "Porsiyon", "Tarifteki Ad", "Malzeme", "Tarifteki Miktar", "Tarifteki Ölçü", "Standart Miktar", "Birim", "Porsiyon Başı",
               "Fire %", "Fiyat Listesinde", "Not"], rr, ai={3, 4, 5, 6, 10}, kod={7, 8, 9}, genis={1: 18, 3: 26, 4: 20, 12: 50})
    for i, r in enumerate(rr, 2):
        if r[10] == "Hayır":
            ws.cell(i, 11).fill = SARI
        if r[6] is None:
            ws.cell(i, 7).fill = KIRMIZI
    g = wb.create_sheet("Hazırlanış ve Alerjen")
    tablo(g, ["Menü Kalemi", "Adımlar", "Alerjenler", "Kodun Bulduğu Alerjenler", "Şefe Sorulacaklar", "Son Kontrol"],
          [[t["ad"], "\n".join(f"{i}. {a}" for i, a in enumerate(sonuc[t["ad"]]["adimlar"], 1)), "\n".join(sonuc[t["ad"]]["alerjenler"]) or "Yok",
            "\n".join(sorted(alerjen_bul(t["metin"]))) or "Yok", "\n".join(sonuc[t["ad"]]["belirsizlikler"]),
            "; ".join(x.split(": ", 1)[1] for x in kalan if x.startswith(t["ad"] + ":")) or "Geçti"] for t in tarifler],
          ai={2, 3, 5}, kod={4}, genis={1: 18, 2: 70, 3: 30, 4: 30, 5: 50, 6: 30})
    k = wb.create_sheet("Kontrol")
    k.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    k.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    k.cell(2, 1).fill = KIRMIZI if kalan else YESIL
    k.append(["Mor: modelin çıkardığı; mavi: kodun dönüşüm tablosuyla hesapladığı. Ev ölçüsü dönüşümleri yaklaşıktır; standart reçeteyi "
              "tartarak doğrulayın. Alerjen listesi tarifteki adlara dayanır; hazır ürünlerin (sos, baharat karışımı, ekmek) etiketindeki "
              "alerjenleri ayrıca kontrol edin. '_receteler.xlsx' ve '_menu.xlsx' 'Porsiyon Maliyeti Hesaplama' kod bloğuna verilebilir."])
    k.column_dimensions["A"].width = 140
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    r = Workbook()
    tablo(r.active, ["Menü Kalemi", "Malzeme", "Miktar", "Birim", "Fire %"], blok, genis={1: 20, 2: 20})
    ry = cikti.with_name(f"{cikti.stem}_receteler.xlsx")
    r.save(ry)
    m = Workbook()
    tablo(m.active, ["Menü Kalemi", "Satış Fiyatı", "Porsiyon Sayısı"], [[t["ad"], None, int(t["porsiyon"]) if t["porsiyon"] else 1] for t in tarifler])
    my = cikti.with_name(f"{cikti.stem}_menu.xlsx")
    m.save(my)
    return ry, my


def calistir(t_yol: Path, f_yol: Path | None, d_yol: Path | None, cikti: Path, tur: int = 2, evet: bool = False) -> dict:
    tarifler, fiyat, donusum = tarifleri_oku(t_yol), fiyat_listesi(f_yol), donusum_oku(d_yol)
    print(f"[OK] {len(tarifler)} tarif · fiyat listesinde {len(fiyat)} malzeme · {len(donusum)} ölçü dönüşümü")
    llm.onay_al(f"{len(tarifler)} tarif metni ve fiyat listesindeki malzeme adları gönderilecek (fiyatlar gönderilmez).", evet)
    sonuc, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(tarifler, fiyat, d, o), lambda x: [s for t in tarifler for s in denetle(t, x[t["ad"]], fiyat)], tur)
    cevrim = {t["ad"]: [cevir(x, donusum, fiyat) for x in sonuc[t["ad"]]["malzemeler"]] for t in tarifler}
    ry, my = rapor_yaz(cikti, tarifler, sonuc, cevrim, fiyat, kalan, yapilan)
    return {"tarifler": tarifler, "sonuc": sonuc, "cevrim": cevrim, "kalan": kalan, "tur": yapilan, "receteler": ry, "menu": my}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Serbest yazılmış tarifleri standart reçete kartına dönüştürür; ölçü çevrimini ve alerjen kontrolünü kod yapar.")
    p.add_argument("--tarifler", type=Path, default=ORNEK / "tarifler.txt", help="Tarifler (.txt/.docx/.pdf); her blok 'Tarif: Ad | Porsiyon: N' ile başlar")
    p.add_argument("--fiyatlar", type=Path, default=ORNEK / "fiyatlar.csv", help="Malzeme; Alış Fiyatı; Birim (Porsiyon Maliyeti kod bloğunun dosyası)")
    p.add_argument("--donusum", type=Path, default=ORNEK / "olcu_donusum.csv", help="Ölçü; Malzeme (* = hepsi); Miktar; Birim (g / ml)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "recete_kartlari.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.tarifler, a.fiyatlar, a.donusum, a.cikti, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    for t in s["tarifler"]:
        c, cv = s["sonuc"][t["ad"]], s["cevrim"][t["ad"]]
        print(f"[OK] {t['ad']} · {len(c['malzemeler'])} malzeme ({sum(x['miktar'] is None for x in cv)} çevrilemedi) · {len(c['adimlar'])} adım · "
              f"alerjen: {', '.join(c['alerjenler']) or 'yok'}")
    for x in s["kalan"]:
        print(f"[!] Kalan: {x}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[OK] Kod bloğu girdileri: {s['receteler'].name}, {s['menu'].name}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
