"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import json
import os
import re
import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
CEVAP = json.loads((Path(__file__).parent / "sahte_cevaplar.json").read_text(encoding="utf-8"))
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    """Mercimek Çorbası'nda ilk turda tereyağına rağmen süt alerjenini yazmaz."""
    MESAJLAR.append(kullanici)
    ad = re.search(r'<tarif ad="([^"]+)"', kullanici)[1]
    return copy.deepcopy(CEVAP[f"{ad}_hatali" if ad == "Mercimek Çorbası" and "<duzeltme>" not in kullanici else ad])


def hazir():
    return ({t["ad"]: t for t in agent.tarifleri_oku(ORNEK / "tarifler.txt")}, agent.fiyat_listesi(ORNEK / "fiyatlar.csv"),
            agent.donusum_oku(ORNEK / "olcu_donusum.csv"))


class KodTesti(unittest.TestCase):
    def test_okuma(self):
        t, f, d = hazir()
        self.assertEqual((t["Mercimek Çorbası"]["porsiyon"], len(t["Mercimek Çorbası"]["satirlar"]), t["Mevsim Salata"]["porsiyon"]), (D(10), 7, D(4)))
        self.assertTrue(t["Mercimek Çorbası"]["hazirlik"].startswith("Soğanları doğrayıp"))
        self.assertEqual(f["tereyagi"], ("Tereyağı", "kg"))
        self.assertEqual((d[("su bardagi", "*")], d[("su bardagi", "kirmizi mercimek")]), ((D(200), "ml"), (D(170), "g")))
        self.assertEqual(agent.metin_sayilari("Yarım limon, 1/2 bardak, 1,5 kg"), {D("0.5"), D("1.5")})

    def test_alerjen(self):
        self.assertEqual(agent.alerjen_bul("3 yemek kaşığı tereyağı, 1 yemek kaşığı un"), {"Gluten içeren tahıllar", "Süt ve süt ürünleri (laktoz dahil)"})
        self.assertEqual(agent.alerjen_bul("tahin, ceviz, karides"), {"Susam", "Sert kabuklu meyveler", "Kabuklular"})
        self.assertEqual(agent.alerjen_bul("domates, salatalık, zeytinyağı"), set())

    def test_denetim(self):
        t, f, _ = hazir()
        self.assertEqual([agent.denetle(t["Mercimek Çorbası"], CEVAP["Mercimek Çorbası"], f), agent.denetle(t["Mevsim Salata"], CEVAP["Mevsim Salata"], f)], [[], []])
        self.assertEqual(agent.denetle(t["Mercimek Çorbası"], CEVAP["Mercimek Çorbası_hatali"], f), ["Mercimek Çorbası: alerjen eksik: Süt ve süt ürünleri (laktoz dahil)"])
        c = copy.deepcopy(CEVAP["Mercimek Çorbası"])
        c["malzemeler"][2]["malzeme"] = "Ayçiçek Yağı"                      # tereyağı → ayçiçek yağı
        c["malzemeler"][0]["miktar"] = "4"
        c["malzemeler"][1]["olcu"] = "kg"
        c["malzemeler"].pop()                                               # limon atlandı
        c["adimlar"][2] = "Mercimeği ekleyip 180 derecede 45 dakika pişirin."
        s = agent.denetle(t["Mercimek Çorbası"], c, f)
        for p in ("'Ayçiçek Yağı' tarifteki adla uyuşmuyor", "miktar '4'", "ölçü 'kg'", "alınmamış: \"Servis için 2 limon\"", "olmayan sayı (süre / sıcaklık): 45, 180"):
            self.assertTrue(any(p in x for x in s), p)

    def test_cevir(self):
        _, f, d = hazir()
        m = lambda **k: {"tarifteki_ad": "", "malzeme": "", "miktar": "", "olcu": "adet", "fire": "", "kanit": ""} | k
        self.assertEqual(agent.cevir(m(malzeme="Kırmızı Mercimek", miktar="5", olcu="su bardağı"), d, f)["miktar"], D(850))
        self.assertEqual(agent.cevir(m(malzeme="Süt", miktar="2", olcu="su bardağı"), d, f)["birim"], "ml")       # * satırı: 400 ml
        self.assertEqual(agent.cevir(m(malzeme="Su", miktar="2", olcu="l"), d, f), {"miktar": D(2000), "birim": "ml", "not": ""})
        self.assertEqual(agent.cevir(m(malzeme="Ekmek", miktar="1", olcu="adet"), d, f)["birim"], "adet")         # fiyat listesi adet
        self.assertIn("dönüşüm yok", agent.cevir(m(malzeme="Sarımsak", miktar="2", olcu="diş"), d, f)["not"])
        self.assertIn("Miktar yok", agent.cevir(m(malzeme="Tuz", olcu="tutam"), d, f)["not"])


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ornek_ve_duzeltme(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "r.xlsx"
            self.assertEqual(agent.main(["--evet", "--cikti", str(cikti)]), 0)
            self.assertEqual(len(MESAJLAR), 3)
            self.assertNotIn("500,00", MESAJLAR[0])                     # fiyatlar gönderilmez
            r = [[c.value for c in x] for x in load_workbook(Path(tmp) / "r_receteler.xlsx").active.iter_rows(min_row=2)]
            self.assertEqual(r[:3], [["Mercimek Çorbası", "Kırmızı Mercimek", 850, "g", None], ["Mercimek Çorbası", "Kuru Soğan", 240, "g", None],
                                     ["Mercimek Çorbası", "Tereyağı", 45, "g", None]])
            self.assertEqual(len(r), 12)
            k = [[c.value for c in x] for x in load_workbook(cikti)["Reçete Kartları"].iter_rows(min_row=2)]
            self.assertEqual((k[0][8], k[4][10]), (85, "Hayır"))         # porsiyon başı 85 g; Su fiyat listesinde yok
            m = [[c.value for c in x] for x in load_workbook(Path(tmp) / "r_menu.xlsx").active.iter_rows(min_row=2)]
            self.assertEqual(m, [["Mercimek Çorbası", None, 10], ["Mevsim Salata", None, 4]])

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(agent.main(["--tarifler", str(Path(tmp) / "yok.txt"), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
