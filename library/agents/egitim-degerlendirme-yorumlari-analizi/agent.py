"""
Eğitim Değerlendirme Yorumları Analizi — Workers / Workless AI Agent
İnsan Kaynakları › Eğitim ve Gelişim Uzmanı

Eğitim sonrası değerlendirme anketlerinin açık uçlu yorumlarını eğitim bazında temalara ayırır, güçlü yönleri ve
geliştirme alanlarını çıkarır, eğitmene geri bildirim notu yazar.
1. Kod anketi okur ("Eğitim Değerlendirme Raporu" kod bloğunun dosyası + yorum sütunları). Katılımcı sütunu gönderilmez;
   katılımcı ve eğitmen adları yorumlarda geçerse takma adla maskelenir. Yorumu --min-yanit'tan az olan eğitim
   anonimlik için analiz edilmez.
2. Kod eğitim bazında yanıt sayısı, memnuniyet ortalaması ve 4–5 verenlerin oranını hesaplar.
3. Model her yorumu tema ve yöne (Olumlu / Olumsuz / Öneri) göre etiketler; güçlü yönleri, geliştirme alanlarını ve
   eğitmen notunu yazar.
4. Kod denetler: her yorum bir kez etiketlenmiş mi; olumsuz / öneri içeren her yorum bir geliştirme alanında mı; düşük
   puanlı (1–2) yanıtın yorumu olumlu sayılmış mı; notta istatistikte olmayan sayı ya da katılımcı adı var mı. Sorun
   varsa yalnız o eğitim modele geri gönderilir. Tema sayımlarını ve sıralamayı kod yapar.

Kullanım:
    python agent.py                                                # örnek 3 eğitim (biri az yanıt nedeniyle atlanır)
    python agent.py --anket degerlendirme.xlsx --min-yanit 5
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import Counter
from datetime import datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import llm

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}
TEMALAR = ["İçerik", "Uygulama ve örnekler", "Eğitmen", "Süre ve tempo", "Materyal", "Ortam ve teknik", "Organizasyon", "Diğer"]
YONLER = ["Olumlu", "Olumsuz", "Öneri"]
YORUM_IZI = ("yorum", "oneri", "begen", "gelistir", "gorus", "eklemek", "aciklama")
SUTUN = {"egitim": ("egitim", "egitim adi", "konu"), "egitmen": ("egitmen", "egitmen adi"), "kisi": ("katilimci", "ad soyad", "personel", "calisan"),
         "puan": ("memnuniyet", "genel memnuniyet", "puan")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def sayi(x) -> Decimal | None:
    s = str(x if x is not None else "").strip().replace(",", ".")
    try:
        return Decimal(s) if s else None
    except InvalidOperation:
        return None


def sayilar(metin: str) -> set[Decimal]:
    return {Decimal(x.replace(",", ".")) for x in re.findall(r"\d+(?:[.,]\d+)?", str(metin))}


def yuvarla(x: Decimal, b: str = "0.1") -> Decimal:
    return x.quantize(Decimal(b), ROUND_HALF_UP)


# ----------------------------------------------------------------------------
# Girdi
# ----------------------------------------------------------------------------

def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: anket dosyası bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, read_only=True, data_only=True)
        s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=metin.split("\n")[0].count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def anket_oku(yol: Path) -> tuple[dict[str, dict], list[str]]:
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: dosya boş.")
    bas = [str(c or "").strip() for c in s[0]]
    kb = [katla(c) for c in bas]
    yer = {k: next((i for i, b in enumerate(kb) if b in es), None) for k, es in SUTUN.items()}
    bilinen = {i for i in yer.values() if i is not None}
    yorum_s = [i for i, b in enumerate(kb) if i not in bilinen and any(x in b for x in YORUM_IZI)]
    if yer["egitim"] is None or not yorum_s:
        raise llm.LLMHatasi(f"{yol.name}: 'Eğitim' sütunu ve en az bir yorum sütunu ('Yorum', 'Beğendiğiniz', 'Geliştirilmesi Gereken' …) gerekli.")
    egitimler, kisiler = {}, []
    for r in s[1:]:
        r = list(r) + [None] * (len(bas) - len(r))
        ad = str(r[yer["egitim"]] or "").strip()
        if not ad:
            continue
        e = egitimler.setdefault(ad, {"ad": ad, "egitmen": "", "puanlar": [], "yanitlar": 0, "yorumlar": []})
        if yer["egitmen"] is not None and r[yer["egitmen"]]:
            e["egitmen"] = str(r[yer["egitmen"]]).strip()
        if yer["kisi"] is not None and r[yer["kisi"]]:
            kisiler.append(str(r[yer["kisi"]]).strip())
        p = sayi(r[yer["puan"]]) if yer["puan"] is not None else None
        yorumlu = False
        for i in yorum_s:
            t = str(r[i] or "").strip()
            if t:
                e["yorumlar"].append({"id": f"Y{len(e['yorumlar']) + 1}", "alan": bas[i], "puan": p, "metin": t})
                yorumlu = True
        if p is not None:
            e["puanlar"].append(p)
        e["yanitlar"] += 1
        e["yorumlu"] = e.get("yorumlu", 0) + yorumlu
    return egitimler, kisiler


def takma_adlar(kisiler: list[str], egitmenler: list[str]) -> dict[str, str]:
    t = {a: f"[EĞİTMEN-{i}]" for i, a in enumerate(dict.fromkeys(x for x in egitmenler if x), 1)}
    for i, a in enumerate(dict.fromkeys(x for x in kisiler if x and x not in t), 1):
        t[a] = f"[KİŞİ-{i}]"
    return t


def maskele(metin: str, adlar: dict[str, str]) -> str:
    for g, t in sorted(adlar.items(), key=lambda x: -len(x[0])):
        metin = re.sub(re.escape(g), t, metin, flags=re.I)
    return llm.maskele(metin)


def istatistik(e: dict, olcek: int) -> dict:
    p = e["puanlar"]
    ort = yuvarla(sum(p) / len(p)) if p else None
    iyi = yuvarla(Decimal(100) * sum(x >= olcek - 1 for x in p) / len(p), "1") if p else None
    return {"yanit": e["yanitlar"], "yorumlu_yanit": e.get("yorumlu", 0), "yorum": len(e["yorumlar"]), "puanli": len(p), "ortalama": ort, "iyi_oran": iyi}


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

GRUP = {"type": "object", "properties": {"baslik": S, "yorumlar": {"type": "array", "items": S}}, "required": ["baslik", "yorumlar"], "additionalProperties": False}
SEMA = {"type": "object", "properties": {
    "etiketler": {"type": "array", "items": {"type": "object", "properties": {
        "id": S, "temalar": {"type": "array", "items": {"type": "object", "properties": {"tema": {"type": "string", "enum": TEMALAR}, "yon": {"type": "string", "enum": YONLER}},
                                                         "required": ["tema", "yon"], "additionalProperties": False}}},
        "required": ["id", "temalar"], "additionalProperties": False}},
    "guclu_yonler": {"type": "array", "items": GRUP},
    "gelistirme_alanlari": {"type": "array", "items": {"type": "object", "properties": {"baslik": S, "yorumlar": {"type": "array", "items": S}, "oneri": S},
                                                     "required": ["baslik", "yorumlar", "oneri"], "additionalProperties": False}},
    "egitmen_notu": S},
    "required": ["etiketler", "guclu_yonler", "gelistirme_alanlari", "egitmen_notu"], "additionalProperties": False}


def mesaj_yap(e: dict, ist: dict, olcek: int) -> str:
    ort = f"{ist['ortalama']}".replace(".", ",") if ist["ortalama"] is not None else "-"
    l = [f"<egitim ad=\"{e['ad']}\" egitmen=\"{e['egitmen_m'] or '-'}\">",
         f"<istatistik>Yanıt: {ist['yanit']} · yorum yazan: {ist['yorumlu_yanit']} · yorum: {ist['yorum']} · memnuniyet ortalaması: {ort} / {olcek}"
         + (f" · {olcek - 1}–{olcek} puan verenler: %{ist['iyi_oran']}" if ist["iyi_oran"] is not None else "") + "</istatistik>"]
    for y in e["yorumlar"]:
        p = f"{y['puan']:g}" if y["puan"] is not None else "-"
        l.append(f"<yorum id=\"{y['id']}\" alan=\"{y['alan']}\" puan=\"{p}\">{y['metin_m']}</yorum>")
    return "\n".join(l + ["</egitim>"])


def denetle(e: dict, ist: dict, c: dict) -> list[str]:
    ad, sorun = e["ad"], []
    yorum = {y["id"]: y for y in e["yorumlar"]}
    say = Counter(x["id"] for x in c["etiketler"])
    eksik = [i for i in yorum if say[i] == 0]
    if eksik:
        sorun.append(f"{ad}: etiketlenmemiş yorum: {', '.join(eksik)}")
    fazla = sorted({i for i in say if i not in yorum} | {i for i, n in say.items() if n > 1 and i in yorum})
    if fazla:
        sorun.append(f"{ad}: bilinmeyen ya da birden çok kez etiketlenen yorum: {', '.join(fazla)}")
    yon = {x["id"]: {t["yon"] for t in x["temalar"]} for x in c["etiketler"]}
    bos = [x["id"] for x in c["etiketler"] if not x["temalar"]]
    if bos:
        sorun.append(f"{ad}: teması boş yorum: {', '.join(bos)}")
    for g in c["guclu_yonler"]:
        yanlis = [i for i in g["yorumlar"] if "Olumlu" not in yon.get(i, set())]
        if yanlis:
            sorun.append(f"{ad}: güçlü yön '{g['baslik']}' olumlu etiketi olmayan yoruma dayanıyor: {', '.join(yanlis)}")
    kapsanan = set()
    for g in c["gelistirme_alanlari"]:
        yanlis = [i for i in g["yorumlar"] if not yon.get(i, set()) & {"Olumsuz", "Öneri"}]
        if yanlis:
            sorun.append(f"{ad}: geliştirme alanı '{g['baslik']}' olumsuz / öneri etiketi olmayan yoruma dayanıyor: {', '.join(yanlis)}")
        if not g["oneri"].strip():
            sorun.append(f"{ad}: geliştirme alanı '{g['baslik']}' için öneri boş")
        kapsanan |= set(g["yorumlar"])
    disarida = [i for i, v in yon.items() if v & {"Olumsuz", "Öneri"} and i in yorum and i not in kapsanan]
    if disarida:
        sorun.append(f"{ad}: olumsuz / öneri içeren yorum hiçbir geliştirme alanında yok: {', '.join(disarida)}")
    dusuk = [i for i, y in yorum.items() if y["puan"] is not None and y["puan"] <= 2 and yon.get(i) == {"Olumlu"}]
    if dusuk:
        sorun.append(f"{ad}: 1–2 puan veren yanıtın yorumu yalnız olumlu etiketlenmiş: {', '.join(dusuk)}; yorumu yeniden okuyun")
    metin = " ".join([c["egitmen_notu"]] + [g["baslik"] for g in c["guclu_yonler"]] + [g["baslik"] + " " + g["oneri"] for g in c["gelistirme_alanlari"]])
    izinli = set().union(*(sayilar(y["metin_m"]) for y in e["yorumlar"])) | {Decimal(ist[k]) for k in ("yanit", "yorumlu_yanit", "yorum") if ist[k]}
    izinli |= {ist[k] for k in ("ortalama", "iyi_oran") if ist[k] is not None} | {Decimal(5), Decimal(4)}
    fazla_s = sorted(sayilar(metin) - izinli)
    if fazla_s:
        sorun.append(f"{ad}: notta istatistikte ya da yorumlarda olmayan sayı: {', '.join(f'{x:g}' for x in fazla_s)} (sayımı kod yapar)")
    if re.search(r"\[KİŞİ-\d+\]", metin):
        sorun.append(f"{ad}: notta katılımcı takma adı var; anket anonimdir, kişi adı yazmayın")
    if not c["egitmen_notu"].strip():
        sorun.append(f"{ad}: eğitmen notu boş")
    return sorun


def uret(egitimler: list[dict], ist: dict, olcek: int, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    sonuc = dict(onceki or {})
    hedef = egitimler if duzeltme is None else [e for e in egitimler if any(s.startswith(e["ad"] + ":") for s in duzeltme)]
    for e in hedef:
        mesaj = mesaj_yap(e, ist[e["ad"]], olcek)
        d = [s.split(": ", 1)[1] for s in (duzeltme or []) if s.startswith(e["ad"] + ":")]
        if d:
            mesaj += ("\n<duzeltme>\nÖnceki cevap: " + json.dumps(sonuc.get(e["ad"]), ensure_ascii=False) + "\nKodun bulduğu sorunlar:\n"
                      + "\n".join(f"- {s}" for s in d) + "\n</duzeltme>")
        sonuc[e["ad"]] = llm.json_iste(sistem, mesaj, SEMA)
    return sonuc


def duzeltme_dongusu(uret_f, denetle_f, tur: int) -> tuple[dict, list[str], int]:
    taslak = uret_f(None, None)
    kalan = denetle_f(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret_f(kalan, taslak)
        kalan = denetle_f(taslak)
    return taslak, kalan, yapilan


def geri_ac(s, adlar: dict[str, str]):
    if isinstance(s, str):
        for g, t in adlar.items():
            s = s.replace(t, g)
        return s
    if isinstance(s, list):
        return [geri_ac(x, adlar) for x in s]
    if isinstance(s, dict):
        return {k: geri_ac(v, adlar) for k, v in s.items()}
    return s


# ----------------------------------------------------------------------------
# Hesap ve çıktı
# ----------------------------------------------------------------------------

def say(e: dict, c: dict) -> dict:
    """Tema × yön sayımı ve geliştirme alanlarının yorum sayısına göre sırası (kod)."""
    tema = Counter((t["tema"], t["yon"]) for x in c["etiketler"] for t in {(z["tema"], z["yon"]): z for z in x["temalar"]}.values())
    yon = Counter(y for x in c["etiketler"] for y in {t["yon"] for t in x["temalar"]})        # yorum sayısı (etiket değil)
    gel = sorted(({**g, "sayi": len(set(g["yorumlar"]))} for g in c["gelistirme_alanlari"]), key=lambda g: -g["sayi"])
    guc = sorted(({**g, "sayi": len(set(g["yorumlar"]))} for g in c["guclu_yonler"]), key=lambda g: -g["sayi"])
    return {"tema": tema, "yon": yon, "gelistirme": gel, "guclu": guc}


def tablo(ws, bas: list[str], satirlar: list[list], ai: set[int] = frozenset(), genis: dict | None = None) -> None:
    ws.append(bas)
    for h in ws[1]:
        h.fill, h.font, h.alignment = BASLIK_DOLGU, BASLIK_YAZI, UST
    for r in satirlar:
        ws.append(r)
        n = ws.max_row
        for j, v in enumerate(r, 1):
            ws.cell(n, j).alignment = UST
            if j in ai and v not in (None, ""):
                ws.cell(n, j).fill = AI_DOLGU
    for j in range(1, len(bas) + 1):
        ws.column_dimensions[get_column_letter(j)].width = (genis or {}).get(j, 14)
    ws.freeze_panes = "A2"


def rapor_yaz(cikti: Path, egitimler: dict, analiz: list[dict], atlanan: list[str], ist: dict, sonuc: dict, sayim: dict, kalan: list[str], tur: int,
              min_yanit: int, olcek: int) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    satir = []
    for ad, e in egitimler.items():
        i = ist[ad]
        if ad in sonuc:
            s = sayim[ad]
            durum = "; ".join(x.split(": ", 1)[1] for x in kalan if x.startswith(ad + ":")) or "Analiz edildi"
            ilk = "\n".join(f"{g['baslik']} ({g['sayi']})" for g in s["gelistirme"][:3])
            satir.append([ad, e["egitmen"], i["yanit"], i["yorum"], float(i["ortalama"]) if i["ortalama"] is not None else None,
                          float(i["iyi_oran"]) if i["iyi_oran"] is not None else None, s["yon"]["Olumlu"], s["yon"]["Olumsuz"], s["yon"]["Öneri"], ilk, durum])
        else:
            satir.append([ad, e["egitmen"], i["yanit"], i["yorum"], float(i["ortalama"]) if i["ortalama"] is not None else None,
                          float(i["iyi_oran"]) if i["iyi_oran"] is not None else None, None, None, None, "",
                          f"Yorum yazan yanıt {i['yorumlu_yanit']} < {min_yanit}: anonimlik için analiz edilmedi"])
    tablo(ws, ["Eğitim", "Eğitmen", "Yanıt", "Yorum", f"Memnuniyet Ort. (/{olcek})", f"{olcek - 1}–{olcek} Oranı (%)", "Olumlu Yorum", "Olumsuz Yorum", "Öneri İçeren Yorum",
               "İlk 3 Geliştirme Alanı (yorum)", "Durum"], satir, ai={10}, genis={1: 22, 2: 18, 10: 44, 11: 40})
    for n, r in enumerate(satir, 2):
        ws.cell(n, 11).fill = YESIL if r[10] == "Analiz edildi" else (SARI if "anonimlik" in r[10] else KIRMIZI)
    g = wb.create_sheet("Geliştirme Alanları")
    rr = []
    for e in analiz:
        y = {x["id"]: x["metin_r"] for x in e["yorumlar"]}
        for a in sayim[e["ad"]]["gelistirme"]:
            rr.append([e["ad"], a["baslik"], a["sayi"], "Tekil yorum" if a["sayi"] == 1 else "", a["oneri"], "\n".join(f"{i}: {y.get(i, '?')}" for i in a["yorumlar"])])
    tablo(g, ["Eğitim", "Geliştirme Alanı", "Yorum Sayısı", "Not", "Öneri", "Dayandığı Yorumlar"], rr, ai={2, 5}, genis={1: 18, 2: 30, 5: 50, 6: 70})
    g = wb.create_sheet("Güçlü Yönler")
    tablo(g, ["Eğitim", "Güçlü Yön", "Yorum Sayısı", "Dayandığı Yorumlar"],
          [[e["ad"], a["baslik"], a["sayi"], ", ".join(a["yorumlar"])] for e in analiz for a in sayim[e["ad"]]["guclu"]], ai={2}, genis={1: 18, 2: 40, 4: 30})
    g = wb.create_sheet("Yorumlar")
    rr = []
    for e in analiz:
        et = {x["id"]: x["temalar"] for x in sonuc[e["ad"]]["etiketler"]}
        for y in e["yorumlar"]:
            rr.append([e["ad"], y["id"], y["alan"], float(y["puan"]) if y["puan"] is not None else None, y["metin_r"],
                       "; ".join(f"{t['tema']} ({t['yon']})" for t in et.get(y["id"], []))])
    tablo(g, ["Eğitim", "ID", "Alan", "Puan", "Yorum", "Tema (Yön)"], rr, ai={6}, genis={1: 18, 3: 22, 5: 70, 6: 40})
    g = wb.create_sheet("Tema Tablosu")
    tablo(g, ["Eğitim", "Tema"] + YONLER, [[e["ad"], t] + [sayim[e["ad"]]["tema"][(t, yn)] for yn in YONLER] for e in analiz for t in TEMALAR
                                          if any(sayim[e["ad"]]["tema"][(t, yn)] for yn in YONLER)], genis={1: 18, 2: 22})
    g = wb.create_sheet("Eğitmen Notları")
    tablo(g, ["Eğitim", "Eğitmen", "Not"], [[e["ad"], e["egitmen"], sonuc[e["ad"]]["egitmen_notu"]] for e in analiz], ai={3}, genis={1: 18, 2: 18, 3: 110})
    k = wb.create_sheet("Kontrol")
    k.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    k.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    k.cell(2, 1).fill = KIRMIZI if kalan else YESIL
    if atlanan:
        k.append([f"Anonimlik için analiz edilmeyen eğitimler (yorum yazan yanıt < {min_yanit}): {', '.join(atlanan)}"])
    k.append(["Mor hücreler modelin etiket ve metinleridir; sayımlar, oranlar ve sıralama koddur. Eğitmen notunu göndermeden önce okuyun. "
              "Katılımcı adı raporda yer almaz; yorumlarda kişiyi tanıtan ayrıntı olabilir, paylaşmadan önce kontrol edin."])
    k.column_dimensions["A"].width = 140
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(yol: Path, cikti: Path, min_yanit: int = 3, olcek: int = 5, tur: int = 2, evet: bool = False) -> dict:
    egitimler, kisiler = anket_oku(yol)
    adlar = takma_adlar(kisiler, [e["egitmen"] for e in egitimler.values()])
    ist = {ad: istatistik(e, olcek) for ad, e in egitimler.items()}
    for e in egitimler.values():
        e["egitmen_m"] = adlar.get(e["egitmen"], e["egitmen"])
        for y in e["yorumlar"]:
            y["metin_m"] = maskele(y["metin"], adlar)
            y["metin_r"] = geri_ac(y["metin_m"], {k: v for k, v in adlar.items() if v.startswith("[EĞİTMEN")})   # raporda katılımcı adı yok
    analiz = [e for ad, e in egitimler.items() if ist[ad]["yorumlu_yanit"] >= min_yanit]
    atlanan = [ad for ad in egitimler if ist[ad]["yorumlu_yanit"] < min_yanit]
    print(f"[OK] {len(egitimler)} eğitim · {sum(i['yorum'] for i in ist.values())} yorum · analiz edilecek {len(analiz)} · az yanıt nedeniyle atlanan {len(atlanan)}")
    sonuc, kalan, yapilan = {}, [], 0
    if analiz:
        llm.onay_al(f"{len(analiz)} eğitimin {sum(len(e['yorumlar']) for e in analiz)} yorumu gönderilecek (katılımcı sütunu gönderilmez; yorumlarda "
                    "geçen katılımcı / eğitmen adları takma adla, telefon ve e-posta maskeli).", evet)
        sonuc, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(analiz, ist, olcek, d, o),
                                                 lambda x: [s for e in analiz for s in denetle(e, ist[e["ad"]], x[e["ad"]])], tur)
        sonuc = geri_ac(sonuc, {k: v for k, v in adlar.items() if v.startswith("[EĞİTMEN")})
        sayim = {e["ad"]: say(e, sonuc[e["ad"]]) for e in analiz}
    else:
        sayim = {}
    rapor_yaz(cikti, egitimler, analiz, atlanan, ist, sonuc, sayim, kalan, yapilan, min_yanit, olcek)
    return {"egitimler": egitimler, "analiz": analiz, "atlanan": atlanan, "ist": ist, "sonuc": sonuc, "sayim": sayim, "kalan": kalan, "tur": yapilan}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Eğitim değerlendirme anketlerinin açık uçlu yorumlarını temalara ayırır ve eğitmen notu yazar.")
    p.add_argument("--anket", type=Path, default=ORNEK / "egitim_degerlendirme.csv", help="Eğitim, Eğitmen, Katılımcı, Memnuniyet + yorum sütunları (.xlsx/.csv)")
    p.add_argument("--min-yanit", type=int, default=3, help="Yorum yazan yanıt bundan azsa eğitim anonimlik için analiz edilmez (varsayılan 3)")
    p.add_argument("--olcek", type=int, default=5, help="Memnuniyet ölçeğinin en yüksek puanı (varsayılan 5)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "egitim_yorum_analizi.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.anket, a.cikti, a.min_yanit, a.olcek, a.tur, a.evet)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    for e in s["analiz"]:
        g = s["sayim"][e["ad"]]["gelistirme"]
        print(f"[OK] {e['ad']} · {len(e['yorumlar'])} yorum · ilk geliştirme alanı: " + (f"{g[0]['baslik']} ({g[0]['sayi']} yorum)" if g else "-"))
    for ad in s["atlanan"]:
        print(f"[i] {ad}: yorum yazan yanıt {s['ist'][ad]['yorumlu_yanit']} < {a.min_yanit}, anonimlik için analiz edilmedi")
    for x in s["kalan"]:
        print(f"[!] Kalan: {x}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
