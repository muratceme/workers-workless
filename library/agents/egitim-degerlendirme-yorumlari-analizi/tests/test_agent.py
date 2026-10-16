"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import json
import os
import re
import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri" / "egitim_degerlendirme.csv"
CEVAP = json.loads((Path(__file__).parent / "sahte_cevaplar.json").read_text(encoding="utf-8"))
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    """İleri Excel'de ilk turda salon / eski Excel yorumunu (Y5) olumsuz etiketler ama hiçbir geliştirme alanına koymaz."""
    MESAJLAR.append(kullanici)
    ad = re.search(r'<egitim ad="([^"]+)"', kullanici)[1]
    return copy.deepcopy(CEVAP[f"{ad}_hatali" if ad == "İleri Excel" and "<duzeltme>" not in kullanici else ad])


def hazirla():
    e, k = agent.anket_oku(ORNEK)
    adlar = agent.takma_adlar(k, [x["egitmen"] for x in e.values()])
    for x in e.values():
        for y in x["yorumlar"]:
            y["metin_m"] = agent.maskele(y["metin"], adlar)
    return e, {a: agent.istatistik(x, 5) for a, x in e.items()}


class KodTesti(unittest.TestCase):
    def test_okuma_ve_istatistik(self):
        e, ist = hazirla()
        self.assertEqual(list(e), ["İleri Excel", "Etkili Sunum", "Zaman Yönetimi"])
        x = e["İleri Excel"]
        self.assertEqual((x["egitmen"], len(x["yorumlar"]), x["yorumlar"][1]["alan"], x["yorumlar"][6]["puan"]), ("Burak Aydın", 9, "Geliştirilmesi Gereken", D(2)))
        self.assertIn("[KİŞİ-1] gibi ileri seviyedekilere", x["yorumlar"][6]["metin_m"])
        self.assertEqual(ist["İleri Excel"], {"yanit": 6, "yorumlu_yanit": 6, "yorum": 9, "puanli": 6, "ortalama": D("3.8"), "iyi_oran": D(67)})
        self.assertEqual(ist["Zaman Yönetimi"]["yorumlu_yanit"], 2)
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "a.csv"
            y.write_text("Eğitim;Memnuniyet\nA;5\n", encoding="utf-8")
            with self.assertRaises(llm.LLMHatasi):
                agent.anket_oku(y)

    def test_denetim(self):
        e, ist = hazirla()
        d = lambda a, c: agent.denetle(e[a], ist[a], c)
        self.assertEqual([d("İleri Excel", CEVAP["İleri Excel"]), d("Etkili Sunum", CEVAP["Etkili Sunum"])], [[], []])
        self.assertEqual(d("İleri Excel", CEVAP["İleri Excel_hatali"]), ["İleri Excel: olumsuz / öneri içeren yorum hiçbir geliştirme alanında yok: Y5"])
        c = copy.deepcopy(CEVAP["Etkili Sunum"])
        c["etiketler"][0]["temalar"] = [{"tema": "Uygulama ve örnekler", "yon": "Olumlu"}]   # 2 puanlı yorum olumlu sayıldı
        c["etiketler"].pop()                                                                # Y7 etiketsiz
        c["etiketler"].append({"id": "Y99", "temalar": []})
        c["guclu_yonler"].append({"baslik": "Uygulama", "yorumlar": ["Y3"]})
        c["egitmen_notu"] += " Katılımcıların %75'i memnun; [KİŞİ-2] uygulama istedi."
        s = d("Etkili Sunum", c)
        for p in ("etiketlenmemiş yorum: Y7", "Y99", "teması boş", "olumlu etiketi olmayan yoruma dayanıyor: Y3",
                  "olumsuz / öneri etiketi olmayan yoruma dayanıyor: Y1", "1–2 puan veren", "75 (sayımı kod yapar)", "takma adı var"):
            self.assertTrue(any(p in x for x in s), p)

    def test_sayim(self):
        e, _ = hazirla()
        s = agent.say(e["İleri Excel"], CEVAP["İleri Excel"])
        self.assertEqual((s["yon"]["Olumlu"], s["yon"]["Olumsuz"], s["yon"]["Öneri"]), (4, 3, 2))   # Y2 iki olumsuz tema ama tek yorum
        self.assertEqual(s["tema"][("Süre ve tempo", "Olumsuz")], 2)
        self.assertEqual([g["sayi"] for g in s["gelistirme"]], [2, 1, 1, 1])


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ornek_ve_duzeltme(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "e.xlsx"
            self.assertEqual(agent.main(["--evet", "--cikti", str(cikti)]), 0)
            self.assertEqual(len(MESAJLAR), 3)                          # 2 eğitim + İleri Excel düzeltmesi; Zaman Yönetimi gönderilmez
            for m in MESAJLAR:
                for ad in ("Deniz Kara", "Burak Aydın", "Zaman Yönetimi"):
                    self.assertNotIn(ad, m)
            wb = load_workbook(cikti)
            o = [[c.value for c in r] for r in wb["Özet"].iter_rows(min_row=2)]
            self.assertEqual(o[0][:9], ["İleri Excel", "Burak Aydın", 6, 9, 3.8, 67, 4, 3, 2])
            self.assertIn("anonimlik", o[2][10])
            g = [[c.value for c in r] for r in wb["Geliştirme Alanları"].iter_rows(min_row=2)]
            self.assertEqual(g[3][1:4], ["Salon sıcaklığı ve eski Excel sürümü", 1, "Tekil yorum"])
            yorumlar = " ".join(str(c.value) for c in wb["Yorumlar"]["E"])
            self.assertNotIn("Deniz Kara", yorumlar)                    # raporda katılımcı adı yok
            self.assertIn("eğitim birimi ele alacak", wb["Eğitmen Notları"]["C2"].value)

    def test_min_yanit_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(ORNEK, Path(tmp) / "e.xlsx", min_yanit=10, evet=True)
            self.assertEqual((s["analiz"], len(s["atlanan"])), ([], 3))
            self.assertEqual(agent.main(["--anket", str(Path(tmp) / "yok.csv"), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
