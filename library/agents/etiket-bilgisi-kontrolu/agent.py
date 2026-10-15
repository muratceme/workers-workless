"""
Etiket Bilgisi Kontrolü — Workers / Workless AI Agent
Kalite Güvence ve Gıda Güvenliği › Gıda Mühendisi

1. Kod etiket metnini okur (ürün adı, bileşenler, net miktar, TETT / SKT, muhafaza, üretici, işletme kayıt no, menşe,
   beslenme bildirimi, ifadeler, temel görüş alanı, x-yüksekliği, en geniş yüzey). Bileşen listesinde vurgulanmış
   alerjenler BÜYÜK HARF ya da **yıldız** ile işaretlenir.
2. Kod Türk Gıda Kodeksi Gıda Etiketleme ve Tüketicileri Bilgilendirme Yönetmeliği'ne göre biçimsel kontrolleri yapar:
   - Zorunlu bilgiler (md. 9) var mı; ad ve net miktar temel görüş alanında mı (md. 16/4)
   - x-yüksekliği en az 1,2 mm mi (en geniş yüzeyi 80 cm²'den küçük ambalajda 0,9 mm) (md. 16)
   - Net miktar birimi (katıda g / kg, sıvıda ml / L) (md. 26); TETT ya da SKT tarihi okunuyor mu
   - Beslenme bildiriminde enerji (kJ ve kcal), yağ, doymuş yağ, karbonhidrat, şeker, protein, tuz bu sırayla var mı
   - "Katkısız" gibi bir ifade varken bileşenlerde katkı maddesi sınıfı / E kodu var mı
   - Anahtar kelimelerle bileşenlerdeki 14 alerjen grubundan olası olanlar
3. Model bileşen listesini yorumlar: alerjen içeren bileşenler (ör. "irmik" → gluten, "peynir altı suyu" → süt), adında
   vurgulanan ve miktarı (%) belirtilmesi gereken bileşenler (QUID), yanıltıcı olabilecek ifadeler.
4. Kod modelin çıktısını denetler (en fazla --tur kez geri gönderilir):
   - Alerjen bileşeni bileşen listesinde aynen geçiyor mu; kodun bulduğu her alerjen grubu modelde var mı
   - QUID bileşeni listede ve ürün adında geçiyor mu; yanıltıcı denen ifade etikette aynen var mı; "katkısız" çelişkisi atlanmış mı
5. Kod sonuçları birleştirir: alerjen vurgulanmış mı, QUID yüzdesi yazılmış mı; kontrol listesini Excel'e yazar.

Kullanım:
    python agent.py                                   # örnek (kurgusal) çilekli kek etiketi
    python agent.py --etiket etiket.docx --tur 3
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
ALANLAR = {"ad": "Ürün Adı", "marka": "Marka", "tur": "Ürün Türü", "tga": "Temel Görüş Alanı", "bilesen": "Bileşenler", "net": "Net Miktar",
           "tett": "TETT", "skt": "SKT", "muhafaza": "Muhafaza Koşulları", "uretici": "Üretici", "kayit": "İşletme Kayıt No", "mense": "Menşe",
           "parti": "Parti No", "beslenme": "Beslenme Bildirimi", "ifadeler": "İfadeler", "alkol": "Alkol", "x": "x-Yüksekliği", "yuzey": "En Geniş Yüzey"}
ZORUNLU = [("ad", "Gıdanın adı"), ("bilesen", "Bileşenler listesi"), ("net", "Net miktar"), ("tarih", "TETT ya da SKT"), ("muhafaza", "Muhafaza / kullanım koşulları"),
           ("uretici", "İşletmecinin adı / unvanı ve adresi"), ("kayit", "İşletme kayıt / onay numarası"), ("beslenme", "Beslenme bildirimi")]
ALERJENLER = ["Gluten içeren tahıllar", "Kabuklular", "Yumurta", "Balık", "Yer fıstığı", "Soya", "Süt", "Sert kabuklu meyveler", "Kereviz", "Hardal",
              "Susam", "Kükürt dioksit ve sülfitler", "Acı bakla (lupin)", "Yumuşakçalar"]
ANAHTAR = {  # kodun kesin tanıdığı kelimeler (model daha fazlasını bulabilir)
    "Gluten içeren tahıllar": r"bugday|cavdar|arpa|yulaf|irmik|gluten|bulgur|kamut|spelt",
    "Kabuklular": r"karides|yengec|istakoz|kerevit|kabuklu deniz",
    "Yumurta": r"yumurta", "Balık": r"\bbalik|ton baligi|hamsi|somon|ançuez|ancuez",
    "Yer fıstığı": r"yer fistig|yerfistig", "Soya": r"\bsoya", "Süt": r"\bsut\b|sut tozu|peynir|tereyag|krema|laktoz|kazein|yogurt|peynir alti",
    "Sert kabuklu meyveler": r"badem|findik|ceviz|kaju|pikan|brezilya|antep fistig|makadamya",
    "Kereviz": r"kereviz", "Hardal": r"hardal", "Susam": r"susam|tahin", "Kükürt dioksit ve sülfitler": r"sulfit|kukurt dioksit|e ?22[0-8]\b",
    "Acı bakla (lupin)": r"aci bakla|lupin", "Yumuşakçalar": r"midye|kalamar|ahtapot|istiridye|salyangoz|yumusakca",
}
KATKI = re.compile(r"\be ?\d{3}[a-z]?\b|jellestirici|asitlik duzenleyici|kabartici|emulgator|koruyucu|renklendirici|tatlandirici|antioksidan|stabilizor|"
                   r"kivam artirici|nem verici|parlatici|topaklanmayi onleyici")
BESLENME = [("Enerji", r"enerji"), ("Yağ", r"\byag\b"), ("Doymuş yağ", r"doymus yag"), ("Karbonhidrat", r"karbonhidrat"), ("Şeker", r"seker"),
            ("Protein", r"protein"), ("Tuz", r"\btuz\b")]

SEMA = {
    "type": "object",
    "properties": {
        "alerjenler": {"type": "array", "items": {"type": "object", "properties": {"kategori": {"type": "string", "enum": ALERJENLER}, "bilesen": S},
                                                  "required": ["kategori", "bilesen"], "additionalProperties": False}},
        "quid": {"type": "array", "items": {"type": "object", "properties": {"bilesen": S, "gerekce": S}, "required": ["bilesen", "gerekce"], "additionalProperties": False}},
        "yaniltici": {"type": "array", "items": {"type": "object", "properties": {"ifade": S, "gerekce": S}, "required": ["ifade", "gerekce"], "additionalProperties": False}},
        "diger_notlar": {"type": "array", "items": S},
    },
    "required": ["alerjenler", "quid", "yaniltici", "diger_notlar"],
    "additionalProperties": False,
}


def etiket_oku(yol: Path) -> dict:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: etiket dosyası bulunamadı.")
    e = {a: "" for a in ALANLAR}
    for s in belge.metin_oku(yol).splitlines():
        a, _, d = s.partition(":")
        k = next((k for k, v in ALANLAR.items() if katla(v) == katla(a)), None)
        if k:
            e[k] = d.strip()
    if not e["ad"] or not e["bilesen"]:
        raise llm.LLMHatasi(f"{yol.name}: 'Ürün Adı:' ve 'Bileşenler:' satırları gerekli.")
    e["bilesenler"] = bilesenleri_ayir(e["bilesen"])
    return e


def bilesenleri_ayir(metin: str) -> list[str]:
    """Bileşen listesini üst ve alt düzey bileşenlere ayırır (parantez içi alt bileşenler de ayrı kayıt olur)."""
    out, kelime = [], ""
    for c in metin.rstrip("."):
        if c in ",()":
            if kelime.strip():
                out.append(kelime.strip())
            kelime = ""
        else:
            kelime += c
    if kelime.strip():
        out.append(kelime.strip())
    return [x for x in out if not re.fullmatch(r"%\s*[\d,.]+", x)]


def vurgulu(bilesen: str, liste: str) -> bool:
    """Bileşen listesinde BÜYÜK HARF ya da **…** ile yazılmış mı?"""
    for m in re.finditer(re.escape(bilesen), liste, flags=re.I):
        parca = liste[m.start():m.end()]
        if parca == parca.upper() and any(ch.isalpha() for ch in parca):
            return True
        if liste[max(0, m.start() - 2):m.start()] == "**" or "**" in liste[max(0, m.start() - 30):m.start()].rsplit(",", 1)[-1]:
            return True
    kelimeler = re.findall(r"\*\*([^*]+)\*\*|\b([A-ZÇĞİÖŞÜ]{3,}(?:\s+[A-ZÇĞİÖŞÜ]{3,})*)\b", liste)
    return any(katla(a or b) and katla(a or b) in katla(bilesen) for a, b in kelimeler)


def kod_alerjenleri(e: dict) -> set[str]:
    k = katla(e["bilesen"])
    return {ad for ad, desen in ANAHTAR.items() if re.search(desen, k)}


def bicimsel_kontrol(e: dict) -> list[dict]:
    """Model gerektirmeyen kontroller: [{kontrol, sonuc (Uygun / Uygunsuz / Kontrol edilmeli), aciklama}]."""
    out = []

    def ekle(kontrol, uygun, aciklama="", belirsiz=False):
        out.append({"kontrol": kontrol, "sonuc": "Kontrol edilmeli" if belirsiz else ("Uygun" if uygun else "Uygunsuz"), "aciklama": aciklama})

    for alan, ad in ZORUNLU:
        var = bool(e["tett"] or e["skt"]) if alan == "tarih" else bool(e[alan])
        ekle(f"Zorunlu bilgi: {ad}", var, "" if var else "Etikette yok")
    tarih = e["skt"] or e["tett"]
    if tarih and not tarih_coz(tarih) and not re.fullmatch(r"\d{2}[./]\d{4}", tarih.strip()):
        ekle("TETT / SKT tarihi", False, f"Tarih okunamadı: '{tarih}' (gün.ay.yıl ya da ay.yıl)")
    if not e["mense"]:
        ekle("Menşe ülke", True, "Yazılmamış; ürünün menşei yanıltıcı olabilecekse zorunludur", belirsiz=True)
    if not e["parti"]:
        ekle("Parti / seri numarası", True, "Yazılmamış; izlenebilirlik için parti numarası gerekir", belirsiz=True)
    tga = katla(e["tga"])
    if e["tga"]:
        for ad in ("Ürün Adı", "Net Miktar"):
            ekle(f"Temel görüş alanı: {ad}", katla(ad) in tga, "" if katla(ad) in tga else "Temel görüş alanında değil (md. 16/4)")
        if e["alkol"]:
            ekle("Temel görüş alanı: Alkol", "alkol" in tga, "" if "alkol" in tga else "Temel görüş alanında değil (md. 16/4)")
    else:
        ekle("Temel görüş alanı", True, "Ön yüzde hangi bilgilerin olduğu belirtilmemiş", belirsiz=True)
    x = re.search(r"[\d.,]+", e["x"])
    yz = re.search(r"[\d.,]+", e["yuzey"])
    if x:
        xv = float(sayi_norm(x.group(0)))
        sinir = 0.9 if yz and float(sayi_norm(yz.group(0))) < 80 else 1.2
        ekle("x-yüksekliği", xv >= sinir, f"{xv:.1f} mm; en az {sinir} mm".replace(".", ",") + " olmalı (md. 16)")
    else:
        ekle("x-yüksekliği", True, "Ölçü verilmemiş; en az 1,2 mm (küçük ambalajda 0,9 mm)", belirsiz=True)
    if e["net"]:
        birim = re.search(r"\b(g|gr|kg|ml|l|lt|cl)\b", katla(e["net"]))
        sivi = katla(e["tur"]) in {"sivi", "icecek"}
        dogru = bool(birim) and ((birim.group(1) in {"ml", "l", "lt", "cl"}) == sivi)
        ekle("Net miktar birimi", dogru, f"'{e['net']}'; " + ("sıvıda hacim (ml, L)" if sivi else "katıda kütle (g, kg)") + " birimi kullanılır (md. 26)")
    if e["beslenme"]:
        b = katla(e["beslenme"])
        yerler = [(ad, m.start() if (m := re.search(d, b)) else -1) for ad, d in BESLENME]
        eksik = [ad for ad, i in yerler if i < 0]
        sira = [i for _, i in yerler if i >= 0]
        if eksik:
            ekle("Beslenme bildirimi öğeleri", False, f"Eksik: {', '.join(eksik)}")
        elif sira != sorted(sira):
            ekle("Beslenme bildirimi sırası", False, "Sıra: enerji, yağ, doymuş yağ, karbonhidrat, şeker, protein, tuz")
        else:
            ekle("Beslenme bildirimi öğeleri ve sırası", True)
        if not ("kj" in b and "kcal" in b):
            ekle("Enerji birimi", False, "Enerji hem kJ hem kcal olarak verilmeli")
    if re.search(r"katkisiz|katki maddesi icermez", katla(e["ifadeler"] + " " + e["ad"])) and KATKI.search(katla(e["bilesen"])):
        ekle("'Katkısız' ifadesi", False, f"Bileşenlerde katkı maddesi var ({KATKI.search(katla(e['bilesen'])).group(0)})")
    return out


def denetle(t: dict, e: dict) -> list[str]:
    sorun = []
    bul = {a["kategori"] for a in t["alerjenler"]}
    for a in t["alerjenler"]:
        if not metinde_var(a["bilesen"], e["bilesen"]):
            sorun.append(f"Alerjen: '{a['bilesen']}' bileşen listesinde aynen geçmiyor")
    for k in sorted(kod_alerjenleri(e) - bul):
        sorun.append(f"Alerjen: bileşenlerde {k} grubundan bileşen var ama listelenmemiş")
    for q in t["quid"]:
        if not metinde_var(q["bilesen"], e["bilesen"]):
            sorun.append(f"QUID: '{q['bilesen']}' bileşen listesinde yok")
        elif not any(w[:4] in katla(e["ad"]) for w in katla(q["bilesen"]).split() if len(w) > 3):
            sorun.append(f"QUID: '{q['bilesen']}' ürün adında vurgulanmıyor; gerekçeyi kontrol edin")
    etiket = e["ifadeler"] + " " + e["ad"] + " " + e["marka"]
    for y in t["yaniltici"]:
        if not metinde_var(y["ifade"], etiket):
            sorun.append(f"Yanıltıcı ifade: '{y['ifade']}' etiketteki ifadelerde yok")
    if any(k["kontrol"] == "'Katkısız' ifadesi" and k["sonuc"] == "Uygunsuz" for k in bicimsel_kontrol(e)) and \
            not any("katkisiz" in katla(y["ifade"]) for y in t["yaniltici"]):
        sorun.append("Yanıltıcı ifade: 'Katkısız' denmiş ama bileşenlerde katkı maddesi var; bu ifadeyi değerlendirin")
    return sorun


def uret(e: dict, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    mesaj = (f"<urun_adi>{e['ad']}</urun_adi>\n<urun_turu>{e['tur'] or '-'}</urun_turu>\n<bilesenler>{e['bilesen']}</bilesenler>\n"
             f"<ifadeler>{e['ifadeler'] or '-'}</ifadeler>\n<beslenme>{e['beslenme'] or '-'}</beslenme>" + duzeltme_blogu(duzeltme, onceki))
    return llm.json_iste(sistem, mesaj, SEMA)


def sonuc_birlestir(t: dict, e: dict) -> list[dict]:
    out = bicimsel_kontrol(e)
    gorulen = set()
    for a in t["alerjenler"]:
        if katla(a["bilesen"]) in gorulen:
            continue
        gorulen.add(katla(a["bilesen"]))
        v = vurgulu(a["bilesen"], e["bilesen"])
        out.append({"kontrol": f"Alerjen vurgusu: {a['bilesen']} ({a['kategori']})", "sonuc": "Uygun" if v else "Uygunsuz",
                    "aciklama": "" if v else "Bileşen listesinde farklı yazı karakteri, stil ya da arka planla vurgulanmalı"})
    for q in t["quid"]:
        m = re.search(re.escape(q["bilesen"]) + r"[^,]{0,25}?%\s*[\d,.]+|%\s*[\d,.]+[^,]{0,5}" + re.escape(q["bilesen"]), e["bilesen"], flags=re.I)
        out.append({"kontrol": f"Bileşen miktarı (QUID): {q['bilesen']}", "sonuc": "Uygun" if m else "Uygunsuz",
                    "aciklama": q["gerekce"] + ("" if m else " — yüzde belirtilmemiş")})
    for y in t["yaniltici"]:
        out.append({"kontrol": f"İfade: {y['ifade']}", "sonuc": "Kontrol edilmeli", "aciklama": y["gerekce"]})
    for n in t["diger_notlar"]:
        out.append({"kontrol": "Not", "sonuc": "Kontrol edilmeli", "aciklama": n})
    return out


def rapor_yaz(cikti: Path, sonuc: list[dict], e: dict, kalan, tur) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kontrol"
    baslik(ws, ["Kontrol", "Sonuç", "Açıklama"], (50, 18, 90))
    for k in sorted(sonuc, key=lambda x: {"Uygunsuz": 0, "Kontrol edilmeli": 1, "Uygun": 2}[x["sonuc"]]):
        ws.append([k["kontrol"], k["sonuc"], k["aciklama"]])
        boya(ws, (2,), {"Uygun": YESIL, "Uygunsuz": KIRMIZI, "Kontrol edilmeli": SARI}[k["sonuc"]])
    son_not(ws, f"Ürün: {e['ad']} · Önemli: Ön kontroldür. Türk Gıda Kodeksi Gıda Etiketleme ve Tüketicileri Bilgilendirme Yönetmeliği, ürüne özgü "
                "kodeks tebliğleri ve güncel kılavuzlar esas alınarak gıda mühendisi tarafından onaylanmalıdır.", kalan, tur)
    hizala(ws)
    b = wb.create_sheet("Bileşenler")
    baslik(b, ["Bileşen", "Vurgulu"], (50, 10))
    for x in e["bilesenler"]:
        b.append([x, "Evet" if vurgulu(x, e["bilesen"]) else ""])
    hizala(b)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(girdi: Path, cikti: Path, tur: int = 2, evet: bool = False) -> dict:
    e = etiket_oku(girdi)
    print(f"[OK] {e['ad']} · {len(e['bilesenler'])} bileşen · kodun bulduğu alerjen grupları: {', '.join(sorted(kod_alerjenleri(e))) or 'yok'}")
    llm.onay_al("Ürün adı, bileşen listesi, ifadeler ve beslenme bildirimi gönderilecek; üretici bilgileri gönderilmez.", evet)
    t, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(e, d, o), lambda x: denetle(x, e), tur)
    sonuc = sonuc_birlestir(t, e)
    rapor_yaz(cikti, sonuc, e, kalan, yapilan)
    return {"etiket": e, "taslak": t, "sonuc": sonuc, "kalan": kalan, "tur": yapilan}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    p = argparse.ArgumentParser(description="Ürün etiketini Türk Gıda Kodeksi etiketleme kurallarına (zorunlu bilgiler, alerjen vurgusu, net miktar) göre kontrol eder.")
    p.add_argument("--etiket", type=Path, default=BURASI / "ornek_veri" / "etiket.txt",
                   help="Etiket metni (.txt/.docx): Ürün Adı, Bileşenler (alerjenler BÜYÜK HARF ya da **…**), Net Miktar, TETT/SKT, … satırları")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "etiket_kontrol.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.etiket, a.cikti, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    uy = [k for k in s["sonuc"] if k["sonuc"] == "Uygunsuz"]
    print(f"[OK] {len(s['sonuc'])} kontrol · {len(uy)} uygunsuz")
    for k in uy:
        print(f"[!] {k['kontrol']}: {k['aciklama']}")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
