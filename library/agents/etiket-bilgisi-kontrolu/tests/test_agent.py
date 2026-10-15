"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ETIKET = KLASOR / "ornek_veri" / "etiket.txt"
MESAJLAR = []


def cevap(dogru=True) -> dict:
    t = {"alerjenler": [{"kategori": "Gluten içeren tahıllar", "bilesen": "Buğday unu"}, {"kategori": "Yumurta", "bilesen": "YUMURTA"},
                        {"kategori": "Süt", "bilesen": "peynir altı suyu tozu"}, {"kategori": "Soya", "bilesen": "soya lesitini"}],
         "quid": [{"bilesen": "çilek dolgusu", "gerekce": "Ürün adında çilek vurgulanıyor."}],
         "yaniltici": [{"ifade": "Katkısız", "gerekce": "Bileşenlerde jelleştirici, asitlik düzenleyici, kabartıcı ve emülgatör var."},
                       {"ifade": "%100 doğal çilek", "gerekce": "Çilek dolgusunda şeker ve glikoz şurubu var; ifade yanıltıcı olabilir."}],
         "diger_notlar": ["Aroma vericinin kaynağı sorulmalı."]}
    if not dogru:
        t["alerjenler"] = [{"kategori": "Gluten içeren tahıllar", "bilesen": "buğday nişastası"}, {"kategori": "Yumurta", "bilesen": "YUMURTA"}]
        t["quid"] = [{"bilesen": "kakao", "gerekce": "Görselde kakao var."}]
        t["yaniltici"] = [{"ifade": "Glutensiz", "gerekce": "Buğday unu var."}]
    return t


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return cevap("<duzeltme>" in kullanici)


class KodTesti(unittest.TestCase):
    def test_bicimsel(self):
        e = agent.etiket_oku(ETIKET)
        self.assertEqual(len(e["bilesenler"]), 20)
        self.assertEqual(agent.kod_alerjenleri(e), {"Gluten içeren tahıllar", "Soya", "Süt", "Yumurta"})
        uy = {k["kontrol"]: k["aciklama"] for k in agent.bicimsel_kontrol(e) if k["sonuc"] == "Uygunsuz"}
        self.assertEqual(set(uy), {"x-yüksekliği", "'Katkısız' ifadesi"})
        self.assertEqual(uy["x-yüksekliği"], "1,0 mm; en az 1,2 mm olmalı (md. 16)")
        self.assertEqual([agent.vurgulu(b, e["bilesen"]) for b in ("YUMURTA", "Buğday unu", "soya lesitini")], [True, False, False])
        self.assertTrue(agent.vurgulu("süt tozu", "şeker, **süt tozu**, kakao"))
        e2 = dict(e, yuzey="60 cm²", x="1,0 mm", net="1 L", tur="Sıvı", beslenme="Enerji 200 kcal; Protein 1 g; Yağ 2 g", tett="", skt="", tga="Marka")
        uy = {k["kontrol"]: k["aciklama"] for k in agent.bicimsel_kontrol(e2) if k["sonuc"] == "Uygunsuz"}
        self.assertNotIn("x-yüksekliği", uy)
        self.assertTrue({"Zorunlu bilgi: TETT ya da SKT", "Beslenme bildirimi öğeleri", "Enerji birimi", "Temel görüş alanı: Ürün Adı"} <= set(uy), uy)

    def test_denetim_ve_sonuc(self):
        e = agent.etiket_oku(ETIKET)
        self.assertEqual(agent.denetle(cevap(), e), [])
        s = agent.denetle(cevap(False), e)
        for parca in ("'buğday nişastası'", "Soya grubundan", "Süt grubundan", "'kakao' bileşen listesinde yok", "'Glutensiz'", "'Katkısız' denmiş"):
            self.assertTrue(any(parca in x for x in s), (parca, s))
        self.assertEqual(len(s), 6, s)
        r = {k["kontrol"]: k["sonuc"] for k in agent.sonuc_birlestir(cevap(), e)}
        self.assertEqual((r["Alerjen vurgusu: YUMURTA (Yumurta)"], r["Alerjen vurgusu: Buğday unu (Gluten içeren tahıllar)"],
                          r["Bileşen miktarı (QUID): çilek dolgusu"]), ("Uygun", "Uygunsuz", "Uygun"))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "e.xlsx"
            s = agent.calistir(ETIKET, cikti, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            self.assertNotIn("Konya", "".join(MESAJLAR))
            ws = load_workbook(cikti)["Kontrol"]
            uygunsuz = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value == "Uygunsuz"]
            self.assertEqual(len(uygunsuz), 5, uygunsuz)
            self.assertEqual(ws["B2"].value, "Uygunsuz")

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(ETIKET, Path(tmp) / "e.xlsx", tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 6)
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            bos = Path(tmp) / "b.txt"
            bos.write_text("Marka: X\n", encoding="utf-8")
            self.assertEqual(agent.main(["--evet", "--etiket", str(bos)]), 1)


if __name__ == "__main__":
    unittest.main()
