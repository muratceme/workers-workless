"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import re
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

RANDEVU = KLASOR / "ornek_veri" / "randevular.csv"
TALIMAT = KLASOR / "ornek_veri" / "hazirlik_talimatlari.csv"
GUN = date(2026, 10, 15)
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    duzelt = "<duzeltme>" in kullanici
    sablonlar = []
    for kimlik, islem, dil, tal in re.findall(r"^(S\d+) \| bölüm: .*? \| işlem: (.*?) \| dil: (.*?) \| onaylı talimatlar: (.*)$", kullanici.split("<duzeltme>")[0], re.M):
        tal = "" if tal == "yok" else tal.replace(" / ", " ")
        if dil == "İngilizce":
            sec = re.findall(r"\d+(?:–\d+)?", tal)
            metin = "Dear {AD}, reminder: {BOLUM}{HEKIM_IFADE} appointment on {TARIH} at {SAAT}. Please follow the preparation: " + ", ".join(sec) + " (see instructions). To change call {ILETISIM}. {KURUM}"
        else:
            metin = "Sayın {AD}, {TARIH} {SAAT} {BOLUM}{HEKIM_IFADE} randevunuz vardır. " + tal + " Değişiklik için {ILETISIM}. {KURUM}"
        if not duzelt:
            if islem == "Gastroskopi":
                metin = metin.replace("8 saat", "6 saat")
            if islem == "Kontrastlı MR" and dil == "Türkçe":
                metin += " Kullandığınız ilaçları bırakın."
            if islem == "Muayene" and "Ortopedi" in kullanici.split(kimlik + " |")[1].split("\n")[0]:
                metin = metin.replace(" Değişiklik için {ILETISIM}.", "")
        sablonlar.append({"kimlik": kimlik, "metin": metin})
    return {"sablonlar": sablonlar}


class KodTesti(unittest.TestCase):
    def test_kombinasyon(self):
        r, notlar = agent.randevu_oku(RANDEVU, GUN)
        k = agent.kombinasyonlar(r, agent.talimat_oku(TALIMAT))
        self.assertEqual((len(r), notlar, len(k)), (9, [], 8))
        self.assertEqual([(x["islem"], x["dil"], len(x["talimatlar"])) for x in k[:4]], [("Kolonoskopi", "TR", 2), ("Gastroskopi", "TR", 2), ("Kontrastlı MR", "TR", 2), ("Kontrastlı MR", "EN", 2)])
        self.assertEqual(k[5]["talimatlar"], ["Randevunuzdan 15 dakika önce kimlik belgenizle kayıt bankosuna gelin."])
        self.assertEqual(agent._sayilar("{SAAT} 8–12 saat, 15 dakika"), {"8", "12", "15"})
        x = {"ad": "", "tarih": GUN, "saat": "09:00", "bolum": "Dahiliye", "hekim": "", "dil": "TR"}
        self.assertEqual(agent.doldur("Sayın {AD}, {BOLUM}{HEKIM_IFADE} {TARIH} {SAAT}  {KURUM}", x, "H", "0"), "Sayın Değerli hastamız, Dahiliye 15.10.2026 09:00 H")


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "r.xlsx"
            s = agent.calistir(RANDEVU, TALIMAT, cikti, GUN, "Örnek Hastanesi", "0212 000 00 00", evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            for m in MESAJLAR:                     # hasta verisi modele gitmez
                self.assertFalse(any(x in m for x in ("Ayşe", "0532", "08:30", "R-90011", "Yılmaz")))
            ilk = MESAJLAR[1].split("<duzeltme>")[1]
            self.assertTrue("talimattaki sayı" in ilk and "'ilac'" in ilk and "{ILETISIM}" in ilk)
            self.assertEqual((len(s["mesajlar"]), len(s["atlanan"])), (8, 1))
            m = s["mesajlar"][0]
            self.assertEqual((m["no"], m["mesaj"][:46]), ("R-90011", "Sayın Ayşe Kaya, 15.10.2026 08:30 Gastroentero"))
            self.assertIn("(Uzm. Dr. A. Yılmaz)", m["mesaj"])
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Şablonlar", "Mesajlar", "Gönderilemeyen"])
            self.assertEqual(len(s["sms"].read_text(encoding="utf-8-sig").splitlines()), 9)

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(RANDEVU, TALIMAT, Path(tmp) / "r.xlsx", GUN, "H", "0", tur=0, evet=True)
            self.assertEqual((len(s["kalan"]), len(s["mesajlar"]), len(s["atlanan"])), (5, 5, 4))
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--tarih", "x"]), 2)
            self.assertEqual(agent.main(["--randevular", str(RANDEVU.with_name("yok.csv")), "--kurum", "H", "--iletisim", "0", "--evet"]), 1)
            kopya = Path(tmp) / "k.csv"
            kopya.write_text(RANDEVU.read_text(encoding="utf-8"), encoding="utf-8")
            self.assertEqual(agent.main(["--randevular", str(kopya), "--tarih", "15.10.2026"]), 2)
            self.assertEqual(agent.main(["--randevular", str(kopya), "--tarih", "01.01.2027", "--kurum", "H", "--iletisim", "0", "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
