"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import json
import os
import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
METIN = (ORNEK / "paketleme_talimati.txt").read_text(encoding="utf-8").strip()
CEVAP = json.loads((Path(__file__).parent / "sahte_cevaplar.json").read_text(encoding="utf-8"))
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    """İlk turda koli ölçüsünü talimattaki 60x40x30 yerine alışılmış 60x40x40 yazar."""
    MESAJLAR.append(kullanici)
    return copy.deepcopy(CEVAP["dogru" if "<duzeltme>" in kullanici else "hatali"])


class KodTesti(unittest.TestCase):
    def test_sevk_ve_yardimcilar(self):
        s = agent.sevk_oku(ORNEK / "sevk.csv")
        self.assertEqual(list(s[0]["bedenler"]), ["S", "M", "L", "XL"])
        self.assertEqual((s[2]["model"], s[2]["agirlik"]), ("SW-02", D("0.55")))
        self.assertEqual(agent.olcu_katla("60 × 40 X 30"), "60x40x30")
        self.assertEqual(len(agent.cumleler(METIN)), 10)

    def test_denetim(self):
        self.assertEqual(agent.denetle(METIN, CEVAP["dogru"]), [])
        self.assertEqual(agent.denetle(METIN, CEVAP["hatali"]), ["koli ölçüsü 60x40x40 talimatta bu hâliyle yok; talimattaki ölçüyü aynen yazın"])
        c = copy.deepcopy(CEVAP["dogru"])
        c["koli_ici_adet"] = "20"                              # talimatta yok, 6'ya bölünmüyor
        c["max_brut_kg"] = ""
        c["asorti"].append({"beden": "XXL", "adet": "1"})
        c["gereklilikler"] = [g for g in c["gereklilikler"] if g["kategori"] != "Belge ve süre"]
        c["gereklilikler"][0]["gereklilik"] += " En fazla 9 renk."
        s = agent.denetle(METIN, c)
        for p in ("koli içi adet '20' talimatta yazmıyor", "tam bölünmüyor", "azami brüt boş", "'XXL' talimatta geçmiyor",
                  "kayda alınmamış: \"Sevkiyattan 5 gün önce", "olmayan sayı: 9"):
            self.assertTrue(any(p in x for x in s), p)
        c = copy.deepcopy(CEVAP["dogru"]) | {"yontem": "solid"}
        self.assertTrue(any("asorti oranı yazılmış" in x for x in agent.denetle(METIN, c)))

    def test_agirlik_ve_komut(self):
        s = agent.sevk_oku(ORNEK / "sevk.csv")
        a = agent.agirlik_kontrolu(CEVAP["dogru"], s, D("0.8"))
        self.assertEqual([(x["model"], x["brut"], x["durum"], x["oneri"]) for x in a],
                         [("TS-01", D("6.08"), "Uygun", None), ("SW-02", D("14.00"), "Sınırı aşıyor", 18)])   # (12-0,8)/0,55 = 20 → 6'nın katı 18
        c = copy.deepcopy(CEVAP["dogru"]) | {"yontem": "solid", "asorti": []}
        self.assertEqual(agent.agirlik_kontrolu(c, s, D("0.8"))[1]["oneri"], 20)
        self.assertEqual(agent.komut(CEVAP["dogru"], ORNEK / "sevk.csv", D("0.8")),
                         "python main.py --sevk sevk.csv --yontem asorti --asorti S=1,M=2,L=2,XL=1 --koli-ici 24 --koli-olcu 60x40x30 --dara 0.8")
        self.assertEqual(agent.beden_uyarilari(CEVAP["dogru"] | {"asorti": [{"beden": "XS", "adet": "1"}]}, s), ["Asorti bedeni 'XS' sevk dosyasında yok."])


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ornek_ve_duzeltme(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "p.xlsx"
            self.assertEqual(agent.main(["--evet", "--cikti", str(cikti)]), 0)
            self.assertEqual(len(MESAJLAR), 2)
            wb = load_workbook(cikti)
            self.assertEqual(wb["Parametreler"]["B5"].value, "60x40x30")
            self.assertEqual(wb["Gereklilikler"].max_row, 7)
            self.assertEqual(wb["Ağırlık Kontrolü"]["F3"].value, "Sınırı aşıyor")
            self.assertIn("en fazla 18", " ".join(str(c.value) for c in wb["Açık Noktalar"]["B"]))

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(agent.main(["--talimat", str(Path(tmp) / "yok.txt"), "--evet"]), 1)
            self.assertEqual(agent.main(["--dara", "abc", "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
