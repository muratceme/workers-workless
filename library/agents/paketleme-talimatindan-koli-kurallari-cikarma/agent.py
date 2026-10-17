"""
Paketleme Talimatından Koli Kuralları Çıkarma — Gerisi Bende AI Agent
Tekstil › Sevkiyat Sorumlusu

Müşterinin paketleme talimatından koli listesi parametrelerini (solid / asorti, asorti oranı, koli içi adet, koli
ölçüsü, azami brüt ağırlık) ve diğer zorunlulukları (polybag, etiket, shipping mark, belge süresi …) çıkarır.
1. Kod talimatı (.txt/.docx/.pdf) ve isteğe bağlı sevk dosyasını ("Koli Listesi (Packing List) Hazırlama" kod bloğunun
   dosyası) okur; telefon, e-posta, TCKN maskelenir.
2. Model parametreleri ve gereklilikleri, talimattan kelimesi kelimesine kanıtlarıyla çıkarır.
3. Kod denetler: sayılar ve koli ölçüsü talimatta aynen var mı; asorti seti koli içine tam bölünüyor mu; talimatta
   ağırlık / ölçü yazarken alan boş mu; kanıtlar talimatta aynen var mı; "-malı / -meli / must / shall" içeren her
   talimat cümlesi bir parametreye ya da gerekliliğe alınmış mı. Sorun varsa model düzeltir.
4. Kod sevk dosyasındaki adet ağırlıklarıyla koli brüt ağırlığını hesaplar; sınırı aşan model için koli içi adet önerir
   (asorti setinin katı olarak) ve kod bloğunun komutunu yazar.

Kullanım:
    python agent.py                                                       # örnek talimat + örnek sevk
    python agent.py --talimat talimat.pdf --sevk sevk.xlsx --dara 0,8
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import re
import sys
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}
EMIR = re.compile(r"(mali|meli)(dir)?\b|\b(zorunlu\w*|gerekir|gereklidir|yasaktir|must|shall|should|required|please|do not|not allowed)\b|(acak|ecek)(tir)?\b")
KATEGORI = ["Polybag", "Etiket", "Koli işareti", "Askı ve katlama", "Belge ve süre", "Koli içeriği", "Diğer"]
SEVK_DISI = {"siparis no", "siparis", "po", "model", "renk", "adet agirligi kg", "adet agirligi", "agirlik", "toplam"}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def sayi(x) -> Decimal | None:
    s = re.sub(r"[^\d,.\-]", "", str(x if x is not None else ""))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return Decimal(s) if s not in ("", "-", ".") else None
    except InvalidOperation:
        return None


def sayilar(metin: str) -> set[Decimal]:
    temiz = re.sub(r"\b[A-Z]{1,4}-\d+\b", " ", str(metin))
    return {v for x in re.findall(r"\d+(?:[.,]\d+)?", temiz) if (v := sayi(x)) is not None}


def olcu_katla(s: str) -> str:
    return re.sub(r"\s*[x×*]\s*", "x", str(s).lower().replace(",", "."))


def cumleler(metin: str) -> list[str]:
    out = []
    for s in metin.splitlines():
        s = re.sub(r"^\s*(\d+[.)]|[-•*])\s*", "", s)
        out += [c.strip() for c in re.split(r"(?<=[.!?;])\s+", s) if c.strip()]
    return out


# ----------------------------------------------------------------------------
# Girdi
# ----------------------------------------------------------------------------

def sevk_oku(yol: Path | None) -> list[dict]:
    if not yol:
        return []
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: sevk dosyası bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, read_only=True, data_only=True)
        s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=metin.split("\n")[0].count)))
    s = [r for r in s if any(c not in (None, "") for c in r)]
    bas = [str(c or "").strip() for c in s[0]]
    kb = [katla(c) for c in bas]
    bedenler = [b for b, k in zip(bas, kb) if k and k not in SEVK_DISI]
    out = []
    for r in s[1:]:
        d = dict(zip(kb, list(r) + [None] * (len(bas) - len(r))))
        out.append({"siparis": str(d.get("siparis no") or d.get("po") or ""), "model": str(d.get("model") or ""), "renk": str(d.get("renk") or ""),
                    "agirlik": sayi(d.get("adet agirligi kg") or d.get("adet agirligi") or d.get("agirlik")),
                    "bedenler": {b: int(sayi(d.get(katla(b))) or 0) for b in bedenler}})
    return out


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

SEMA = {"type": "object", "properties": {
    "yontem": {"type": "string", "enum": ["asorti", "solid", "belirsiz"]},
    "asorti": {"type": "array", "items": {"type": "object", "properties": {"beden": S, "adet": S}, "required": ["beden", "adet"], "additionalProperties": False}},
    "koli_ici_adet": S, "koli_olcu": S, "max_brut_kg": S,
    "kanitlar": {"type": "array", "items": {"type": "object", "properties": {"alan": {"type": "string", "enum": ["yontem", "koli_ici_adet", "koli_olcu", "max_brut_kg"]},
                                                                            "kanit": S}, "required": ["alan", "kanit"], "additionalProperties": False}},
    "gereklilikler": {"type": "array", "items": {"type": "object", "properties": {"kategori": {"type": "string", "enum": KATEGORI}, "gereklilik": S, "kanit": S},
                                                 "required": ["kategori", "gereklilik", "kanit"], "additionalProperties": False}},
    "belirsizlikler": {"type": "array", "items": S}},
    "required": ["yontem", "asorti", "koli_ici_adet", "koli_olcu", "max_brut_kg", "kanitlar", "gereklilikler", "belirsizlikler"], "additionalProperties": False}


def set_toplami(c: dict) -> int | None:
    try:
        return sum(int(sayi(a["adet"])) for a in c["asorti"]) if c["asorti"] else None
    except TypeError:
        return None


def denetle(metin: str, c: dict) -> list[str]:
    km, sorun = katla(metin), []
    ns = sayilar(metin)
    for ad, v in (("koli içi adet", c["koli_ici_adet"]), ("azami brüt", c["max_brut_kg"]), *((f"asorti {a['beden']}", a["adet"]) for a in c["asorti"])):
        if v.strip() and sayi(v) not in ns:
            sorun.append(f"{ad} '{v}' talimatta yazmıyor")
    if c["koli_olcu"].strip():
        o = olcu_katla(re.sub(r"\s*cm\s*$", "", c["koli_olcu"].strip(), flags=re.I))
        if not re.fullmatch(r"\d+(\.\d+)?x\d+(\.\d+)?x\d+(\.\d+)?", o):
            sorun.append(f"koli ölçüsü 'UxGxY' biçiminde olmalı ('{c['koli_olcu']}')")
        elif o not in olcu_katla(metin):
            sorun.append(f"koli ölçüsü {c['koli_olcu']} talimatta bu hâliyle yok; talimattaki ölçüyü aynen yazın")
    elif re.search(r"\d+\s*[x×*]\s*\d+\s*[x×*]\s*\d+", metin):
        sorun.append("talimatta koli ölçüsü var ama alan boş")
    if not c["max_brut_kg"].strip() and re.search(r"\d+(?:[.,]\d+)?\s*kg", metin, re.I):
        sorun.append("talimatta kg cinsinden sınır var ama azami brüt boş")
    if c["yontem"] == "asorti":
        st = set_toplami(c)
        if not st:
            sorun.append("asorti seçildi ama beden oranı yok ya da okunamadı")
        else:
            k = sayi(c["koli_ici_adet"])
            if k and k % st:
                sorun.append(f"asorti seti {st} adet; koli içi {k:g} buna tam bölünmüyor")
            for a in c["asorti"]:
                if not re.search(rf"(?<![a-z0-9]){re.escape(katla(a['beden']))}(?![a-z0-9])", km):
                    sorun.append(f"asorti bedeni '{a['beden']}' talimatta geçmiyor")
    elif c["asorti"]:
        sorun.append(f"yöntem '{c['yontem']}' iken asorti oranı yazılmış")
    kanitlar = []
    for x in c["kanitlar"] + [{"alan": g["kategori"], "kanit": g["kanit"]} for g in c["gereklilikler"]]:
        kk = katla(x["kanit"])
        if len(kk.split()) < 2 or kk not in km:
            sorun.append(f"kanıt talimatta aynen geçmiyor ({x['alan']}): \"{x['kanit'][:80]}\"")
        else:
            kanitlar.append(kk)
    alanlar = {x["alan"] for x in c["kanitlar"]}
    for ad, v in (("yontem", c["yontem"] != "belirsiz"), ("koli_ici_adet", c["koli_ici_adet"]), ("koli_olcu", c["koli_olcu"]), ("max_brut_kg", c["max_brut_kg"])):
        if v and ad not in alanlar:
            sorun.append(f"'{ad}' için kanıt yok")
    fazla = sorted(set().union(*(sayilar(g["gereklilik"]) for g in c["gereklilikler"])) - ns) if c["gereklilikler"] else []
    if fazla:
        sorun.append(f"gerekliliklerde talimatta olmayan sayı: {', '.join(f'{x:g}' for x in fazla)}")
    for cu in cumleler(metin):
        kc = katla(cu)
        if len(kc.split()) >= 4 and EMIR.search(kc) and not any(x in kc or kc in x for x in kanitlar):
            sorun.append(f"talimat cümlesi kayda alınmamış: \"{cu[:110]}\"")
    return sorun


def uret(metin: str, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    mesaj = f"<talimat>\n{metin}\n</talimat>"
    if duzeltme:
        mesaj += ("\n<duzeltme>\nÖnceki cevap: " + json.dumps(onceki, ensure_ascii=False) + "\nKodun bulduğu sorunlar:\n"
                  + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")
    return llm.json_iste(sistem, mesaj, SEMA)


def duzeltme_dongusu(uret_f, denetle_f, tur: int) -> tuple[dict, list[str], int]:
    taslak = uret_f(None, None)
    kalan = denetle_f(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret_f(kalan, taslak)
        kalan = denetle_f(taslak)
    return taslak, kalan, yapilan


# ----------------------------------------------------------------------------
# Hesap ve çıktı
# ----------------------------------------------------------------------------

def agirlik_kontrolu(c: dict, sevk: list[dict], dara: Decimal) -> list[dict]:
    k, mx, st = sayi(c["koli_ici_adet"]), sayi(c["max_brut_kg"]), set_toplami(c) if c["yontem"] == "asorti" else 1
    out, gorulen = [], set()
    for r in sevk:
        anah = (r["model"], r["agirlik"])
        if not r["agirlik"] or anah in gorulen:
            continue
        gorulen.add(anah)
        brut = k * r["agirlik"] + dara if k else None
        durum, oneri = "", None
        if brut is not None and mx:
            if brut <= mx:
                durum = "Uygun"
            else:
                en_fazla = math.floor((mx - dara) / r["agirlik"])
                oneri = en_fazla - en_fazla % (st or 1)
                durum = "Sınırı aşıyor"
        out.append({"model": r["model"], "agirlik": r["agirlik"], "koli_ici": k, "brut": brut, "max": mx, "durum": durum or "Sınır yok", "oneri": oneri})
    return out


def beden_uyarilari(c: dict, sevk: list[dict]) -> list[str]:
    if c["yontem"] != "asorti" or not sevk:
        return []
    sb = {katla(b) for r in sevk for b in r["bedenler"]}
    return [f"Asorti bedeni '{a['beden']}' sevk dosyasında yok." for a in c["asorti"] if katla(a["beden"]) not in sb]


def komut(c: dict, sevk_yolu: Path | None, dara: Decimal) -> str:
    p = ["python main.py", f"--sevk {sevk_yolu.name if sevk_yolu else 'sevk.xlsx'}"]
    if c["yontem"] in ("asorti", "solid"):
        p.append(f"--yontem {c['yontem']}")
    if c["yontem"] == "asorti" and c["asorti"]:
        p.append("--asorti " + ",".join(f"{a['beden']}={int(sayi(a['adet']))}" for a in c["asorti"]))
    if c["koli_ici_adet"]:
        p.append(f"--koli-ici {int(sayi(c['koli_ici_adet']))}")
    if c["koli_olcu"]:
        p.append("--koli-olcu " + olcu_katla(re.sub(r"\s*cm\s*$", "", c["koli_olcu"].strip(), flags=re.I)))
    p.append(f"--dara {dara:g}")
    return " ".join(p)


def tablo(ws, bas: list[str], satirlar: list[list], ai: set[int] = frozenset(), genis: dict | None = None) -> None:
    ws.append(bas)
    for h in ws[1]:
        h.fill, h.font, h.alignment = BASLIK_DOLGU, BASLIK_YAZI, UST
    for r in satirlar:
        ws.append(r)
        for j, v in enumerate(r, 1):
            ws.cell(ws.max_row, j).alignment = UST
            if j in ai and v not in (None, ""):
                ws.cell(ws.max_row, j).fill = AI_DOLGU
    for j in range(1, len(bas) + 1):
        ws.column_dimensions[get_column_letter(j)].width = (genis or {}).get(j, 16)
    ws.freeze_panes = "A2"


def rapor_yaz(cikti: Path, c: dict, ak: list[dict], uyarilar: list[str], kmt: str, kalan: list[str], tur: int) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Parametreler"
    kn = {x["alan"]: x["kanit"] for x in c["kanitlar"]}
    tablo(ws, ["Parametre", "Değer", "Talimattaki Kanıt"],
          [["Yöntem", c["yontem"], kn.get("yontem", "")], ["Asorti", ", ".join(f"{a['beden']}={a['adet']}" for a in c["asorti"]), kn.get("yontem", "")],
           ["Koli içi adet", c["koli_ici_adet"], kn.get("koli_ici_adet", "")], ["Koli ölçüsü (cm)", c["koli_olcu"], kn.get("koli_olcu", "")],
           ["Azami brüt (kg)", c["max_brut_kg"], kn.get("max_brut_kg", "")], [], ["Kod bloğu komutu", kmt, ""]], ai={2, 3}, genis={1: 20, 2: 50, 3: 80})
    g = wb.create_sheet("Gereklilikler")
    tablo(g, ["No", "Kategori", "Gereklilik", "Talimattaki Kanıt", "Kontrol Edildi", "Kontrol Eden"],
          [[i, x["kategori"], x["gereklilik"], x["kanit"], "", ""] for i, x in enumerate(c["gereklilikler"], 1)], ai={2, 3, 4}, genis={1: 5, 3: 50, 4: 70})
    g = wb.create_sheet("Ağırlık Kontrolü")
    tablo(g, ["Model", "Adet Ağırlığı (kg)", "Koli İçi", "Koli Brüt (kg)", "Azami Brüt (kg)", "Durum", "Önerilen Koli İçi"],
          [[x["model"], float(x["agirlik"]), int(x["koli_ici"]) if x["koli_ici"] else None, float(x["brut"]) if x["brut"] is not None else None,
            float(x["max"]) if x["max"] else None, x["durum"], x["oneri"]] for x in ak])
    for i, x in enumerate(ak, 2):
        g.cell(i, 6).fill = KIRMIZI if x["durum"] == "Sınırı aşıyor" else (YESIL if x["durum"] == "Uygun" else PatternFill())
    g = wb.create_sheet("Açık Noktalar")
    tablo(g, ["Kaynak", "Açıklama"], [["Model", b] for b in c["belirsizlikler"]] + [["Kod", u] for u in uyarilar]
          + [["Kod", f"{x['model']}: {int(x['koli_ici'])} adetlik koli {x['brut']:.2f} kg, sınır {x['max']:g} kg; koli içi en fazla {x['oneri']} olabilir. "
                     "Müşteriden onay alın (koli içi adet ve ağırlık sınırı birlikte sağlanamıyor)."] for x in ak if x["durum"] == "Sınırı aşıyor"],
          genis={2: 110})
    k = wb.create_sheet("Kontrol")
    k.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    k.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    k.cell(2, 1).fill = KIRMIZI if kalan else YESIL
    k.append(["Mor hücreler modelin çıkardığı değerlerdir; ağırlık hesabı ve öneriler koddur. Gereklilikler listesini sevkiyat öncesi kontrol "
              "listesi olarak kullanın. Brüt ağırlık = koli içi × adet ağırlığı + koli darası (--dara)."])
    k.column_dimensions["A"].width = 140
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(t_yol: Path, sevk_yolu: Path | None, cikti: Path, dara: Decimal = Decimal("0.8"), tur: int = 2, evet: bool = False) -> dict:
    if not t_yol.exists():
        raise llm.LLMHatasi(f"{t_yol}: talimat dosyası bulunamadı.")
    metin = llm.maskele(belge.metin_oku(t_yol).replace("\r", "").strip())
    if not metin:
        raise llm.LLMHatasi(f"{t_yol.name}: talimat boş.")
    sevk = sevk_oku(sevk_yolu)
    print(f"[OK] Talimat {len(cumleler(metin))} cümle · sevk satırı {len(sevk)}")
    llm.onay_al("Paketleme talimatının metni gönderilecek (telefon, e-posta, TCKN maskeli).", evet)
    c, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(metin, d, o), lambda x: denetle(metin, x), tur)
    ak = agirlik_kontrolu(c, sevk, dara)
    uy = beden_uyarilari(c, sevk)
    kmt = komut(c, sevk_yolu, dara)
    rapor_yaz(cikti, c, ak, uy, kmt, kalan, yapilan)
    return {"cevap": c, "agirlik": ak, "uyarilar": uy, "komut": kmt, "kalan": kalan, "tur": yapilan}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Müşteri paketleme talimatından koli kurallarını ve gereklilikleri çıkarır; ağırlık sınırını kod kontrol eder.")
    p.add_argument("--talimat", type=Path, default=ORNEK / "paketleme_talimati.txt", help="Paketleme talimatı (.txt/.docx/.pdf)")
    p.add_argument("--sevk", type=Path, help="Sevk dosyası (Koli Listesi kod bloğunun dosyası: Sipariş No, Model, Renk, bedenler, Adet Ağırlığı)")
    p.add_argument("--dara", default="0,8", help="Boş koli ağırlığı, kg (varsayılan 0,8)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "paketleme_kurallari.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    sevk = a.sevk or (ORNEK / "sevk.csv" if a.talimat == ORNEK / "paketleme_talimati.txt" else None)
    try:
        dara = sayi(a.dara)
        if dara is None or dara < 0:
            raise llm.LLMHatasi(f"--dara okunamadı: '{a.dara}'")
        s = calistir(a.talimat, sevk, a.cikti, dara, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    c = s["cevap"]
    print(f"[OK] {c['yontem']} · koli içi {c['koli_ici_adet'] or '-'} · ölçü {c['koli_olcu'] or '-'} · azami brüt {c['max_brut_kg'] or '-'} kg · "
          f"{len(c['gereklilikler'])} gereklilik")
    for x in s["agirlik"]:
        if x["durum"] == "Sınırı aşıyor":
            print(f"[!] {x['model']}: koli {x['brut']:.2f} kg > {x['max']:g} kg → koli içi en fazla {x['oneri']}")
    for x in s["uyarilar"] + [f"Kalan: {k}" for k in s["kalan"]]:
        print(f"[!] {x}")
    print(f"[OK] Kod bloğu komutu: {s['komut']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
