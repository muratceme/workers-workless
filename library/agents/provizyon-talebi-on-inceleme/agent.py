"""
Provizyon Talebi Ön İnceleme — Gerisi Bende AI Agent
Sağlık Sigortası › Provizyon Uzmanı (Anlaşmalı Kurumlar Uzmanı da kullanır)

Anlaşmalı kurumlardan gelen provizyon taleplerini (form, epikriz, hekim notu) poliçe bilgileri ve verdiğiniz özel / genel
şartlar metniyle ön incelemeden geçirir. Karar vermez; karar provizyon uzmanının ve gerektiğinde hekimindir.
1. Kod talep belgelerini, poliçe tablosunu ve şartlar metnini okur. "Hasta:" satırındaki ad ve --gizle ile verilen adlar,
   TCKN, telefon, e-posta sağlayıcıya gönderilmez.
2. Model talebi çıkarır (poliçe no, yatış türü ve tarihi, gün, tanı ve ICD-10 kodu, işlem, talep tutarı), poliçe öncesi
   var olabilecek hastalık işaretlerini ve ilgili şart maddelerini belgeden alıntıyla gösterir, eksik bilgileri ve bir
   öneri (Onay önerilir / Tıbbi inceleme gerekir / Ek bilgi iste) yazar.
3. Kod çıkarımı denetler: alıntılar belgede aynen geçiyor mu; madde şartlar metninde var mı; tarih, tutar ve ICD kodu
   belgede yazıyor mu; poliçe tabloda var mı. Sorun varsa yalnız o talep modele geri gönderilir (en fazla --tur kez).
4. Kod poliçe kontrollerini yapar: planlı yatışta bekleme süresi (başlangıç + bekleme > yatış tarihi), poliçe başlangıcından
   önceki işlem, yatarak / ayakta limit, katılım payı ve sigorta payı. Kodun bulguları ya da modelin işaretleri varsa
   durum "İnceleme" olur.

Kullanım:
    python agent.py                                                    # örnek 3 talep
    python agent.py --talepler ./talepler --policeler policeler.xlsx --sartlar ozel_sartlar.txt
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
RENK = {"Uygun": "E3F5E1", "İnceleme": "FFF4CE", "Eksik bilgi": "FDE2E1"}
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}
MADDE = re.compile(r"\b(?:madde|md\.?)\s*(\d+(?:\.\d+)*)", re.I)
ICD = re.compile(r"^[A-TV-Z]\d{2}(?:\.\d{1,2})?$")
HASTA_SATIRI = re.compile(r"^(\s*(?:hasta|hastanın|sigortalı)(?: adı)?(?: soyadı)?\s*:\s*)(.+)$", re.I | re.M)


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1 or (s.count(".") == 1 and len(s.split(".")[1]) == 3):
        s = s.replace(".", "")
    try:
        return Decimal(s) if s not in ("", "-", ".") else None
    except InvalidOperation:
        return None


def sayilar(metin: str) -> set[Decimal]:
    temiz = re.sub(r"\b\d{1,2}\.\d{1,2}\.\d{4}\b|\b\d{1,2}:\d{2}\b|\b[A-Z]\d{2}\.\d{1,2}\b", " ", str(metin))
    return {v for x in re.findall(r"\d+(?:[.,]\d+)*", temiz) if (v := para(x)) is not None}


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    m = re.search(r"\b(\d{1,2})\.(\d{1,2})\.(\d{4})\b", str(x or ""))
    if not m:
        return None
    try:
        return date(int(m[3]), int(m[2]), int(m[1]))
    except ValueError:
        return None


def tl(x: Decimal) -> str:
    return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


# ----------------------------------------------------------------------------
# Girdi
# ----------------------------------------------------------------------------

POLICE = {"no": ("police no", "police"), "baslangic": ("baslangic tarihi", "baslangic", "police baslangic"), "plan": ("plan", "urun"),
          "yatarak": ("yatarak limit", "yatarak tedavi limiti"), "ayakta": ("ayakta limit", "ayakta tedavi limiti"), "katilim": ("katilim %", "katilim payi", "katilim"),
          "bekleme": ("bekleme suresi (gun)", "bekleme suresi", "bekleme")}


def tablo(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, read_only=True, data_only=True)
        s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=metin.split("\n")[0].count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def policeleri_oku(yol: Path) -> dict[str, dict]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: poliçe tablosu bulunamadı.")
    s = tablo(yol)
    b = [katla(c) for c in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in POLICE.items()}
    if idx["no"] is None or idx["baslangic"] is None:
        raise llm.LLMHatasi(f"{yol.name}: 'Poliçe No' ve 'Başlangıç Tarihi' sütunları gerekli.")
    out = {}
    for r in s[1:]:
        al = lambda a: None if idx[a] is None or idx[a] >= len(r) else r[idx[a]]  # noqa: E731
        if al("no"):
            out[str(al("no")).strip()] = {"baslangic": tarih_coz(al("baslangic")), "plan": str(al("plan") or ""), "yatarak": para(al("yatarak")),
                                          "ayakta": para(al("ayakta")), "katilim": para(al("katilim")) or Decimal(0),
                                          "bekleme": int(para(al("bekleme")) or 0)}
    return out


def takma_adlar(adlar: list[str]) -> dict[str, str]:
    return {a: f"[HASTA-{i}]" for i, a in enumerate(dict.fromkeys(x.strip() for x in adlar if x.strip()), 1)}


def maskele(metin: str, kisiler: dict[str, str]) -> tuple[str, dict[str, str]]:
    """'Hasta: Ad Soyad' satırındaki ad ve --gizle adları [HASTA-n] yapılır; ortak maskeler uygulanır."""
    k = dict(kisiler)
    for m in HASTA_SATIRI.finditer(metin):
        ad = m[2].strip()
        if ad and not ad.startswith("[") and ad not in k:
            k[ad] = f"[HASTA-{len(k) + 1}]"
    for g, t in sorted(k.items(), key=lambda x: -len(x[0])):
        metin = re.sub(re.escape(g), t, metin, flags=re.I)
    return llm.maskele(metin), k


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

ALINTI = {"type": "object", "properties": {"alinti": S, "aciklama": S}, "required": ["alinti", "aciklama"], "additionalProperties": False}
MADDE_S = {"type": "object", "properties": {"madde": S, "alinti": S, "gerekce": S}, "required": ["madde", "alinti", "gerekce"], "additionalProperties": False}
SEMA = {"type": "object", "properties": {
    "police_no": S, "talep_tarihi": S, "islem_tarihi": S, "yatis_turu": {"type": "string", "enum": ["Planlı yatış", "Acil yatış", "Ayakta"]},
    "yatis_gun": S, "tani": S, "icd": S, "islem": S, "talep_tutari": S, "onceden_var_isaretleri": {"type": "array", "items": ALINTI},
    "ilgili_maddeler": {"type": "array", "items": MADDE_S}, "eksik_bilgiler": {"type": "array", "items": S},
    "oneri": {"type": "string", "enum": ["Onay önerilir", "Tıbbi inceleme gerekir", "Ek bilgi iste"]}, "oneri_gerekcesi": S},
    "required": ["police_no", "talep_tarihi", "islem_tarihi", "yatis_turu", "yatis_gun", "tani", "icd", "islem", "talep_tutari",
                 "onceden_var_isaretleri", "ilgili_maddeler", "eksik_bilgiler", "oneri", "oneri_gerekcesi"], "additionalProperties": False}


def denetle(t: dict, c: dict, sartlar: str, policeler: dict) -> list[str]:
    on, m, sorun = t["no"], t["maskeli"], []
    for a, ad in (("tani", "Tanı"), ("islem", "İşlem"), ("police_no", "Poliçe no")):
        if c[a].strip() and not metinde_var(c[a], m):
            sorun.append(f"{on}: {ad} '{c[a][:50]}' belgede geçmiyor")
    if c["police_no"].strip() and c["police_no"].strip() not in policeler:
        sorun.append(f"{on}: poliçe '{c['police_no']}' poliçe tablosunda yok")
    for a, ad in (("talep_tarihi", "Talep tarihi"), ("islem_tarihi", "İşlem / yatış tarihi")):
        if c[a].strip():
            d = tarih_coz(c[a])
            if d is None:
                sorun.append(f"{on}: {ad} '{c[a]}' GG.AA.YYYY değil")
            elif f"{d:%d.%m.%Y}" not in m:
                sorun.append(f"{on}: {ad} {d:%d.%m.%Y} belgede yazmıyor")
    for a, ad in (("talep_tutari", "Talep tutarı"), ("yatis_gun", "Yatış günü")):
        if c[a].strip() and (para(c[a]) is None or para(c[a]) not in sayilar(m)):
            sorun.append(f"{on}: {ad} '{c[a]}' belgede yazmıyor")
    if c["icd"].strip():
        if not ICD.match(c["icd"].strip()):
            sorun.append(f"{on}: ICD-10 kodu '{c['icd']}' biçim dışı (ör. K80.2)")
        elif c["icd"].strip() not in m:
            sorun.append(f"{on}: ICD-10 kodu {c['icd']} belgede yazmıyor; kod uydurmayın, yoksa boş bırakın")
    for i, x in enumerate(c["onceden_var_isaretleri"], 1):
        if not metinde_var(x["alinti"], m):
            sorun.append(f"{on}: önceden var işareti {i}: alıntı '{x['alinti'][:50]}' belgede aynen geçmiyor")
    gecerli = {g.lower() for g in MADDE.findall(sartlar)}
    for i, x in enumerate(c["ilgili_maddeler"], 1):
        g = MADDE.search(x["madde"])
        if not g or g[1].lower() not in gecerli:
            sorun.append(f"{on}: ilgili madde {i} '{x['madde']}' şartlar metninde yok")
        elif not metinde_var(x["alinti"], sartlar):
            sorun.append(f"{on}: ilgili madde {i}: alıntı şartlar metninde aynen geçmiyor")
    if c["oneri"] == "Onay önerilir" and c["onceden_var_isaretleri"]:
        sorun.append(f"{on}: önceden var olabilecek hastalık işareti varken 'Onay önerilir' seçilmiş; 'Tıbbi inceleme gerekir' olmalı")
    if re.search(r"\[(HASTA-\d+|TCKN|TELEFON)\]", c["oneri_gerekcesi"] + " ".join(x["aciklama"] for x in c["onceden_var_isaretleri"])):
        sorun.append(f"{on}: açıklamada hasta yer tutucusu var; 'sigortalı' deyin")
    return sorun


def police_kontrol(c: dict, p: dict | None) -> dict:
    """Kod: bekleme süresi, başlangıç, limit, katılım ve sigorta payı."""
    b, h = [], {"katilim": None, "sigorta_payi": None, "limit": None}
    if p is None:
        return {"bulgu": [("Hata", "Poliçe tabloda yok: poliçe numarasını kontrol edin")], **h}
    islem, tutar = tarih_coz(c["islem_tarihi"]), para(c["talep_tutari"])
    if islem is None:
        b.append(("Dikkat", "İşlem / yatış tarihi yok: bekleme süresi kontrol edilemedi"))
    elif p["baslangic"] and islem < p["baslangic"]:
        b.append(("Hata", f"İşlem tarihi ({islem:%d.%m.%Y}) poliçe başlangıcından ({p['baslangic']:%d.%m.%Y}) önce"))
    elif p["baslangic"] and c["yatis_turu"] == "Planlı yatış" and p["bekleme"]:
        bitis = p["baslangic"] + timedelta(days=p["bekleme"])
        if islem < bitis:
            b.append(("Dikkat", f"Planlı yatış bekleme süresi içinde: başlangıç {p['baslangic']:%d.%m.%Y} + {p['bekleme']} gün = {bitis:%d.%m.%Y}, "
                                f"yatış {islem:%d.%m.%Y} ({(islem - p['baslangic']).days}. gün)"))
    limit = p["ayakta"] if c["yatis_turu"] == "Ayakta" else p["yatarak"]
    h["limit"] = limit
    if tutar is None:
        b.append(("Dikkat", "Talep tutarı yok"))
    else:
        esas = min(tutar, limit) if limit is not None else tutar
        if limit is not None and tutar > limit:
            b.append(("Dikkat", f"Talep {tl(tutar)} TL, {'ayakta' if c['yatis_turu'] == 'Ayakta' else 'yatarak'} limit {tl(limit)} TL"))
        h["katilim"] = (esas * p["katilim"] / 100).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        h["sigorta_payi"] = esas - h["katilim"]
    return {"bulgu": b, **h}


def uret(talepler: list[dict], sartlar: str, policeler: dict, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    sonuc = dict(onceki or {})
    hedef = talepler if duzeltme is None else [t for t in talepler if any(s.startswith(t["no"] + ":") for s in duzeltme)]
    for t in hedef:
        bilinen = next((no for no in policeler if no in t["maskeli"]), None)            # belgede geçen poliçenin başlangıcı modele verilir
        police = (f"<police no=\"{bilinen}\" baslangic=\"{policeler[bilinen]['baslangic']:%d.%m.%Y}\" plan=\"{policeler[bilinen]['plan']}\"/>\n"
                  if bilinen and policeler[bilinen]["baslangic"] else "<police>(poliçe tabloda bulunamadı)</police>\n")
        mesaj = (f"<sartlar>\n{sartlar or '(şartlar metni verilmedi: ilgili_maddeler boş kalsın)'}\n</sartlar>\n{police}"
                 f"<talep dosya=\"{t['dosya']}\">\n{t['maskeli']}\n</talep>")
        d = [s.split(": ", 1)[1] for s in (duzeltme or []) if s.startswith(t["no"] + ":")]
        if d:
            mesaj += ("\n<duzeltme>\nÖnceki inceleme: " + json.dumps(sonuc.get(t["no"]), ensure_ascii=False) + "\nKodun bulduğu sorunlar:\n"
                      + "\n".join(f"- {s}" for s in d) + "\n</duzeltme>")
        sonuc[t["no"]] = llm.json_iste(sistem, mesaj, SEMA)
    return sonuc


def duzeltme_dongusu(uret_f, denetle_f, tur: int) -> tuple[dict, list[str], int]:
    taslak = uret_f(None, None)
    kalan = denetle_f(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret_f(kalan, taslak)
        kalan = denetle_f(taslak)
    return taslak, kalan, yapilan


# ----------------------------------------------------------------------------
# Çıktı
# ----------------------------------------------------------------------------

def _baslik(ws, basliklar, genislik) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font, h.alignment = BASLIK_DOLGU, BASLIK_YAZI, UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def _f(x):
    return None if x is None else float(x)


def rapor_yaz(cikti: Path, talepler: list[dict], sonuc: dict, kalan: list[str], tur: int) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Talepler"
    _baslik(ws, ["Talep", "Poliçe No", "Yatış Türü", "Yatış Tarihi", "Gün", "Tanı", "ICD-10", "İşlem", "Talep Tutarı", "Limit", "Katılım Payı",
                 "Sigorta Payı (ön hesap)", "Kod Bulguları", "Önceden Var Olabilecek Hastalık İşaretleri", "İlgili Şart Maddeleri", "Eksik Bilgiler",
                 "Model Önerisi", "Öneri Gerekçesi", "Durum", "Son Kontrol"],
            (14, 14, 12, 11, 5, 26, 8, 30, 12, 12, 11, 13, 50, 50, 50, 40, 18, 40, 11, 30))
    for t in talepler:
        c, k = sonuc[t["no"]], t["kontrol"]
        sor = [x.split(": ", 1)[1] for x in kalan if x.startswith(t["no"] + ":")]
        isaret = "\n".join(f"“{x['alinti']}” — {x['aciklama']}" for x in c["onceden_var_isaretleri"])
        maddeler = "\n".join(f"{x['madde']}: {x['gerekce']}" for x in c["ilgili_maddeler"])
        ws.append([t["dosya"].rsplit(".", 1)[0], c["police_no"], c["yatis_turu"], tarih_coz(c["islem_tarihi"]), _f(para(c["yatis_gun"])), c["tani"], c["icd"],
                   c["islem"], _f(para(c["talep_tutari"])), _f(k["limit"]), _f(k["katilim"]), _f(k["sigorta_payi"]),
                   "\n".join(f"{o}: {m}" for o, m in k["bulgu"]), isaret, maddeler, "\n".join(f"• {x}" for x in c["eksik_bilgiler"]), c["oneri"],
                   c["oneri_gerekcesi"], t["durum"], "; ".join(sor) or "Geçti"])
        n = ws.max_row
        ws.cell(n, 4).number_format = "DD.MM.YYYY"
        for j in (9, 10, 11, 12):
            ws.cell(n, j).number_format = "#,##0.00"
        for j in (3, 6, 7, 8, 14, 15, 16, 17, 18):
            ws.cell(n, j).fill = AI_DOLGU
        ws.cell(n, 19).fill = PatternFill("solid", fgColor=RENK[t["durum"]])
        ws.cell(n, 20).fill = KIRMIZI if sor else YESIL
        for h in ws[n]:
            h.alignment = UST
    k = wb.create_sheet("Kontrol")
    k.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    k.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    k.cell(2, 1).fill = KIRMIZI if kalan else YESIL
    k.append(["Önemli: Bu bir ön incelemedir; provizyon onayı veya reddi değildir. Önceden var olabilecek hastalık işaretleri yalnız belgedeki "
              "ifadelerdir, tıbbi değerlendirme yerine geçmez; ret kararı hekim ve provizyon uzmanı tarafından poliçe şartlarına göre verilir. "
              "Katılım ve sigorta payı ön hesaptır; anlaşma fiyatı ve fatura kontrolü ayrıca yapılmalıdır."])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(girdi: Path, cikti: Path, police_yolu: Path, sart_yolu: Path | None = None, gizle: list[str] | None = None, tur: int = 2,
             evet: bool = False) -> dict:
    policeler = policeleri_oku(police_yolu)
    if sart_yolu is not None and not sart_yolu.exists():
        raise llm.LLMHatasi(f"{sart_yolu}: şartlar metni bulunamadı.")
    sartlar = belge.metin_oku(sart_yolu) if sart_yolu else ""
    if girdi.is_dir():
        dosyalar = sorted(p for p in girdi.iterdir() if p.is_file() and p.suffix.lower() in belge.DESTEKLENEN)
    elif girdi.exists():
        dosyalar = [girdi]
    else:
        raise belge.BelgeHatasi(f"{girdi}: bulunamadı.")
    if not dosyalar:
        raise belge.BelgeHatasi(f"{girdi}: talep belgesi yok.")
    kisiler = takma_adlar(gizle or [])
    talepler = []
    for i, p in enumerate(dosyalar, 1):
        m, kisiler = maskele(belge.metin_oku(p), kisiler)
        talepler.append({"no": f"T{i}", "dosya": p.name, "maskeli": m})
    print(f"[OK] {len(talepler)} talep · poliçe: {len(policeler)} · şartlar: {len(sartlar):,} karakter")
    llm.onay_al(f"{len(talepler)} provizyon talebinin tanı, işlem ve hekim notları gönderilecek (sağlık verisi). Hasta adları ('Hasta:' satırı ve "
                "--gizle), TCKN, telefon, e-posta maskeli; mümkünse yerel model (Ollama) kullanın.", evet)
    sonuc, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(talepler, sartlar, policeler, d, o),
                                             lambda x: [s for t in talepler for s in denetle(t, x[t["no"]], sartlar, policeler)], tur)
    for t in talepler:
        c = sonuc[t["no"]]
        t["kontrol"] = police_kontrol(c, policeler.get(c["police_no"].strip()))
        onem = {o for o, _ in t["kontrol"]["bulgu"]}
        t["durum"] = ("Eksik bilgi" if "Hata" in onem or c["oneri"] == "Ek bilgi iste" else
                      "İnceleme" if onem or c["onceden_var_isaretleri"] or c["oneri"] != "Onay önerilir" else "Uygun")
    rapor_yaz(cikti, talepler, sonuc, kalan, yapilan)
    return {"talepler": talepler, "sonuc": sonuc, "kalan": kalan, "tur": yapilan}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Provizyon taleplerini poliçe bilgileri ve şartlar metniyle ön incelemeden geçirir (karar vermez).")
    p.add_argument("--talepler", type=Path, default=ORNEK / "talepler", help="Talep belgesi ya da klasörü (.pdf, .docx, .txt)")
    p.add_argument("--policeler", type=Path, default=ORNEK / "policeler.csv", help="Poliçe tablosu (.xlsx/.csv): Poliçe No, Başlangıç Tarihi, Yatarak / Ayakta Limit, Katılım %%, Bekleme Süresi (gün)")
    p.add_argument("--sartlar", type=Path, default=None, help="Özel / genel şartlar metni (.txt/.docx/.pdf)")
    p.add_argument("--gizle", nargs="*", default=[], metavar="AD", help="Ayrıca maskelenecek kişi adları")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "provizyon_on_inceleme.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    sartlar = a.sartlar or (ORNEK / "genel_sartlar.txt" if a.talepler == ORNEK / "talepler" else None)
    try:
        s = calistir(a.talepler, a.cikti, a.policeler, sartlar, a.gizle, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    for t in s["talepler"]:
        c = s["sonuc"][t["no"]]
        print(f"[OK] {t['dosya']} · {c['police_no']} · {c['tani'][:30]} · {t['durum']}")
        for o, m in t["kontrol"]["bulgu"]:
            print(f"[!]   {o}: {m}")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
