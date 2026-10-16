"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import json
import os
import re
import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
CEVAP = json.loads((Path(__file__).parent / "sahte_cevaplar.json").read_text(encoding="utf-8"))
SARTLAR = (ORNEK / "genel_sartlar.txt").read_text(encoding="utf-8")
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    """PRV-1002'de önceden var işaretleri varken ilk turda 'Onay önerilir' der."""
    MESAJLAR.append(kullanici)
    d = re.search(r'<talep dosya="([^"]+)\.txt"', kullanici)[1]
    if d == "PRV-1002" and "<duzeltme>" not in kullanici:
        return copy.deepcopy(CEVAP["PRV-1002_hatali"])
    return copy.deepcopy(CEVAP[d])


def talepler():
    out, k = [], {}
    for i, p in enumerate(sorted((ORNEK / "talepler").iterdir()), 1):
        m, k = agent.maskele(p.read_text(encoding="utf-8"), k)
        out.append({"no": f"T{i}", "dosya": p.name, "maskeli": m})
    return out, k


class KodTesti(unittest.TestCase):
    def test_okuma_ve_maskeleme(self):
        p = agent.policeleri_oku(ORNEK / "policeler.csv")
        self.assertEqual({k: (v["baslangic"], v["bekleme"], v["katilim"], v["yatarak"]) for k, v in p.items()}["TSS-26-0457"],
                         (date(2026, 9, 20), 90, D(20), D("150000.00")))
        t, k = talepler()
        self.assertEqual(k, {"Ayşe Deneme": "[HASTA-1]", "Can Örnek": "[HASTA-2]", "Zeynep Kurgu": "[HASTA-3]"})
        for x in t:
            self.assertNotRegex(x["maskeli"], r"Ayşe|Can Örnek|Zeynep")
        m, _ = agent.maskele("Hasta Adı Soyadı: Ali Veli\nAli Veli 3 gündür ateşli. TCKN 10000000146", {})
        self.assertEqual(m, "Hasta Adı Soyadı: [HASTA-1]\n[HASTA-1] 3 gündür ateşli. TCKN [TCKN]")

    def test_police_kontrol(self):
        p = agent.policeleri_oku(ORNEK / "policeler.csv")
        k = agent.police_kontrol(CEVAP["PRV-1002"], p["TSS-26-0457"])
        self.assertEqual((k["katilim"], k["sigorta_payi"], k["limit"]), (D("18400.00"), D("73600.00"), D("150000.00")))
        self.assertEqual(k["bulgu"], [("Dikkat", "Planlı yatış bekleme süresi içinde: başlangıç 20.09.2026 + 90 gün = 19.12.2026, yatış 20.10.2026 (30. gün)")])
        self.assertEqual(agent.police_kontrol(CEVAP["PRV-1003"] | {"islem_tarihi": "15.10.2026"}, p["TSS-26-0457"])["bulgu"], [])   # acil: bekleme yok
        k = agent.police_kontrol(CEVAP["PRV-1001"] | {"talep_tutari": "260.000,00"}, p["TSS-26-0101"])
        self.assertEqual((k["bulgu"][0][0], k["katilim"], k["sigorta_payi"]), ("Dikkat", D("25000.00"), D("225000.00")))         # limit üstü: limitten hesap
        self.assertEqual(agent.police_kontrol(CEVAP["PRV-1001"] | {"islem_tarihi": "15.01.2026"}, p["TSS-26-0101"])["bulgu"][0][0], "Hata")
        self.assertEqual(agent.police_kontrol(CEVAP["PRV-1001"], None)["bulgu"][0][0], "Hata")

    def test_denetim(self):
        t, _ = talepler()
        p = agent.policeleri_oku(ORNEK / "policeler.csv")
        self.assertEqual([agent.denetle(x, CEVAP[x["dosya"][:-4]], SARTLAR, p) for x in t], [[], [], []])
        self.assertIn("'Onay önerilir' seçilmiş", agent.denetle(t[1], CEVAP["PRV-1002_hatali"], SARTLAR, p)[0])
        c = copy.deepcopy(CEVAP["PRV-1002"])
        c["icd"] = "M51.2"
        c["onceden_var_isaretleri"][0]["alinti"] = "5 yıldır bel ağrısı"
        c["ilgili_maddeler"][0]["madde"] = "Madde 3.9"
        c["talep_tutari"] = "95.000,00"
        c["police_no"] = "TSS-26-9999"
        s = agent.denetle(t[1], c, SARTLAR, p)
        for parca in ("ICD-10 kodu M51.2 belgede yazmıyor", "alıntı '5 yıldır", "'Madde 3.9' şartlar metninde yok", "Talep tutarı '95.000,00'",
                      "Poliçe no 'TSS-26-9999' belgede geçmiyor"):
            self.assertTrue(any(parca in x for x in s), parca)
        self.assertIn("biçim dışı", agent.denetle(t[0], CEVAP["PRV-1001"] | {"icd": "kolelitiazis"}, SARTLAR, p)[0])


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ornek_ve_duzeltme(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "p.xlsx"
            s = agent.calistir(ORNEK / "talepler", cikti, ORNEK / "policeler.csv", ORNEK / "genel_sartlar.txt", evet=True)
            self.assertEqual((len(MESAJLAR), s["tur"], s["kalan"]), (4, 1, []))
            self.assertIn('<police no="TSS-26-0101" baslangic="01.02.2026" plan="Gold"/>', MESAJLAR[0])
            for m in MESAJLAR:
                self.assertNotRegex(m, r"Ayşe|Can Örnek|Zeynep")
            self.assertEqual([t["durum"] for t in s["talepler"]], ["Uygun", "İnceleme", "Uygun"])
            ws = load_workbook(cikti)["Talepler"]
            self.assertEqual([ws.cell(3, j).value for j in (1, 7, 9, 10, 11, 12, 17, 19)],
                             ["PRV-1002", "M51.1", 92000, 150000, 18400, 73600, "Tıbbi inceleme gerekir", "İnceleme"])
            self.assertIn("bekleme süresi içinde", ws.cell(3, 13).value)
            self.assertIn("“4 yıldır bel ağrısı ve sol bacakta uyuşma”", ws.cell(3, 14).value)

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "p.xlsx")]), 0)
            self.assertEqual(agent.main(["--talepler", str(Path(tmp) / "yok"), "--evet"]), 1)
            self.assertEqual(agent.main(["--policeler", str(Path(tmp) / "yok.csv"), "--evet"]), 1)
            self.assertEqual(agent.main(["--sartlar", str(Path(tmp) / "yok.txt"), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
