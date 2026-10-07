"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("WW_PROVIDER", "anthropic")
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
GONDERILENLER = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    GONDERILENLER.append(kullanici)
    puan = 85 if "İşe alım" in kullanici or "işe alım" in kullanici.lower() else 30
    return {
        "son_pozisyon": "Uzman", "son_sirket": "Örnek A.Ş.", "toplam_deneyim_yil": 5,
        "en_yuksek_egitim": "Lisans", "bolum": "İşletme", "diller": ["İngilizce"],
        "beceriler": ["Excel"], "ozet": "Kısa özet.", "ilan_uyum_puani": puan,
        "guclu_yonler": ["a"], "eksikler": ["b"], "gerekce": "c",
    }


class MaskelemeTesti(unittest.TestCase):
    def test_kisisel_veriler_maskelenir(self):
        m = llm.maskele("ali@x.com 0532 111 22 33 TR33 0006 1005 1978 6457 8413 26 12345678901")
        for etiket in ("[E-POSTA]", "[TELEFON]", "[IBAN]", "[TCKN]"):
            self.assertIn(etiket, m)
        self.assertNotIn("ali@x.com", m)


class AgentTesti(unittest.TestCase):
    def setUp(self):
        GONDERILENLER.clear()

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_uctan_uca_ornek_veri(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "r.xlsx"
            cvler = agent.calistir(ORNEK / "cvler", cikti, ORNEK / "ilan.txt", evet=True, paralel=2)

            self.assertEqual(len(cvler), 4)
            self.assertEqual(len(GONDERILENLER), 3)  # bozuk dosya gönderilmez
            hepsi = "\n".join(GONDERILENLER)
            for gizli in ("ayse.yilmaz@ornek-mail.com", "0532 111 22 33", "Ayşe Yılmaz", "MEHMET KAYA"):
                self.assertNotIn(gizli, hepsi)
            self.assertIn("<is_ilani>", hepsi)

            ws = load_workbook(cikti)["Adaylar"]
            basliklar = [h.value for h in ws[1]]
            self.assertEqual(basliklar[-1], "İnsan Kararı")
            self.assertEqual(ws.cell(2, 2).value, "Ayşe Yılmaz")       # en yüksek puan ilk sırada
            self.assertEqual(ws.cell(2, 3).value, "ayse.yilmaz@ornek-mail.com")  # iletişim yerelde eklenir

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ilan_yoksa_puan_bos(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            cvler = agent.calistir(ORNEK / "cvler", Path(tmp) / "r.xlsx", None, evet=True)
            self.assertTrue(all(c.profil.get("ilan_uyum_puani") is None for c in cvler if c.profil))

    def test_onaysiz_etkilesimsiz_calisma_durur(self):
        with mock.patch.object(sys, "stdin", None):
            with self.assertRaises(llm.LLMHatasi):
                llm.onay_al("deneme", evet=False)


if __name__ == "__main__":
    unittest.main()
