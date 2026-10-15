"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

BILGI = KLASOR / "ornek_veri" / "kampanya.txt"
MESAJLAR = []


def taslak(dogru=True) -> dict:
    t = {"konu_satirlari": ["Sonbahar ürünlerinde %30'a varan indirim başladı", "Sonbahar İndirimi: seçili ürünlerde %30'a varan fırsat",
                            "Gardırobunu yenile: %30'a varan sonbahar indirimi", "SONBAHAR10 ile sepette ek %10 indirim",
                            "Sonbahar İndirimi 26 Ekim'e kadar sürüyor"],
         "on_baslik": "Seçili sonbahar ürünlerinde indirim 26 Ekim'e kadar.",
         "baslik": "Sonbahar İndirimi başladı",
         "govde": ["Merhaba,", "Seçili sonbahar ürünlerinde %30'a varan indirim seni bekliyor. Kampanya 20.10.2026–26.10.2026 tarihleri arasında "
                               "online mağazada geçerli.",
                   "1.500 TL ve üzeri alışverişlerinde SONBAHAR10 koduyla sepette ek %10 indirim seni bekliyor. 750 TL ve üzeri siparişlerde kargo ücretsiz."],
         "buton": "Fırsatları keşfet",
         "kosul_ozeti": "İndirim oranları ürüne göre değişir. Kod üye başına bir kez kullanılabilir. Kampanya 26.10.2026'da sona erer."}
    if not dogru:
        t["konu_satirlari"][0] = "SONBAHAR FIRSATI: %30 indirim!!"
        t["govde"][1] = t["govde"][1].replace("Seçili sonbahar ürünlerinde", "Tüm ürünlerde")
        t["govde"][2] = "750 TL ve üzeri siparişlerde kargo ücretsiz."
        t["kosul_ozeti"] += " Detaylar: https://ornekmagaza.example/kosullar"
    return t


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return taslak("<duzeltme>" in kullanici)


class KodTesti(unittest.TestCase):
    def test_okuma(self):
        b = agent.bilgiyi_oku(BILGI)
        self.assertEqual((b["kod"], b["bitis"], len(b["kosullar"])), ("SONBAHAR10", "26.10.2026", 4))
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "k.txt"
            y.write_text(BILGI.read_text(encoding="utf-8").replace("MERSİS No: 0123456789000015\n", ""), encoding="utf-8")
            with self.assertRaises(llm.LLMHatasi):
                agent.bilgiyi_oku(y)

    def test_denetim(self):
        b = agent.bilgiyi_oku(BILGI)
        self.assertEqual(agent.denetle(taslak(), b, 60, 5), [])
        s = agent.denetle(taslak(False), b, 60, 5)
        self.assertEqual(len(s), 6, s)
        self.assertTrue(any("büyük harf" in x for x in s) and any("ünlem" in x for x in s) and any("yanıltıcı" in x for x in s)
                        and any("tum urunler" in x for x in s) and any("SONBAHAR10" in x for x in s) and any("bağlantı" in x for x in s))
        self.assertTrue(any("sınır 40" in x for x in agent.denetle(taslak(), b, 40, 5)))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "e.xlsx"
            s = agent.calistir(BILGI, cikti, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            for gizli in ("0123456789000015", "musteri@", "abonelik-iptal"):
                self.assertNotIn(gizli, MESAJLAR[0])
            h = cikti.with_suffix(".html").read_text(encoding="utf-8")
            self.assertIn('href="https://ornekmagaza.example/sonbahar"', h)
            self.assertIn("MERSİS No: 0123456789000015", h)
            self.assertIn('href="https://ornekmagaza.example/abonelik-iptal"', h)
            self.assertIn("Abonelikten çık: https://ornekmagaza.example/abonelik-iptal", cikti.with_suffix(".txt").read_text(encoding="utf-8"))
            self.assertEqual(load_workbook(cikti).sheetnames, ["E-bülten"])

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(BILGI, Path(tmp) / "e.xlsx", tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 6)
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--evet", "--girdi", str(BILGI.with_name("yok.txt"))]), 1)


if __name__ == "__main__":
    unittest.main()
