"""
E-bülten Metni Hazırlama — Gerisi Bende AI Agent
Pazarlama › Dijital Pazarlama Uzmanı

1. Kod kampanya bilgisini okur (kampanya, tarihler, indirim, kod, koşullar, bağlantı) ve ticari elektronik ileti için
   gönderen bilgilerini (unvan, MERSİS no, iletişim, ret bağlantısı) denetler.
2. Model konu satırı alternatifleri, ön başlık, başlık, gövde, buton metni ve koşul özeti yazar.
3. Kod denetler:
   - Konu satırları sayısı ve uzunluğu; tamamı büyük harf kelime, çoklu ünlem ve spam ifadeleri
   - Metindeki her sayı (yüzde, tarih, fiyat) kampanya bilgisinde var mı
   - Bilgide "%30'a varan" yazıyorsa metinde "varan / kadar" düşürülmüş mü (yanıltıcı indirim)
   - Bitiş tarihi ve kampanya kodu metinde var mı
   - "Ücretsiz", "hediye", "garanti", "en ucuz" gibi iddialar bilgide geçiyor mu; metne bağlantı (URL) yazılmış mı
   Sorun varsa modele geri gönderilir (en fazla --tur kez).
4. Kod; kampanya bağlantısını butona, gönderen bilgilerini ve ret (abonelikten çıkma) bağlantısını alt bilgiye ekleyip
   HTML ve düz metin e-bülten taslağı yazar.

Kullanım:
    python agent.py                                        # örnek sonbahar kampanyası
    python agent.py --girdi kampanya.txt --konu-azami 50
"""
from __future__ import annotations

import argparse
import csv
import html
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
SEMA = {
    "type": "object",
    "properties": {
        "konu_satirlari": {"type": "array", "items": S}, "on_baslik": S, "baslik": S,
        "govde": {"type": "array", "items": S}, "buton": S, "kosul_ozeti": S,
    },
    "required": ["konu_satirlari", "on_baslik", "baslik", "govde", "buton", "kosul_ozeti"],
    "additionalProperties": False,
}
ALANLAR = {"unvan": "Gönderen Unvan", "mersis": "MERSİS No", "iletisim": "İletişim", "ret": "Ret Bağlantısı", "kampanya": "Kampanya",
           "baslangic": "Başlangıç", "bitis": "Bitiş", "indirim": "İndirim", "kod": "Kod", "baglanti": "Bağlantı", "hedef": "Hedef Kitle", "ton": "Ton"}
ZORUNLU = ("unvan", "mersis", "iletisim", "ret", "kampanya", "bitis", "baglanti")
SPAM = re.compile(r"\b(bedava|hemen kazan\w*|tikla kazan\w*|100 garanti\w*|kacirma\w*\s+firsat\w*|son sans\w*|para kazan\w*|risk(?:siz|sız))\b")
IDDIA = re.compile(r"\b(ucretsiz|bedava|hediye\w*|garanti\w*|en ucuz|en dusuk fiyat\w*|en uygun fiyat\w*|stoklarla sinirli|sinirli sayida|tum urunler\w*)\b")
URL = re.compile(r"https?://|www\.", re.I)
AYLAR = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]
EK_YUZDE = r"(?:'?\s*(?:[ae]|y[ae]|ye|ya)?\s*(?:varan|kadar))"


def bilgiyi_oku(yol: Path) -> dict:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: kampanya bilgi dosyası bulunamadı.")
    b, kosullar, bolum = {a: "" for a in ALANLAR}, [], None
    for s in belge.metin_oku(yol).splitlines():
        if not s.strip():
            continue
        a, _, d = s.partition(":")
        anahtar = next((k for k, v in ALANLAR.items() if katla(v) == katla(a)), None)
        if anahtar:
            b[anahtar], bolum = d.strip(), None
        elif katla(s) == "kosullar":
            bolum = "kosullar"
        elif bolum == "kosullar":
            kosullar.append(s.strip().lstrip("-•* ").strip())
    b["kosullar"] = kosullar
    eksik = [ALANLAR[a] for a in ZORUNLU if not b[a]]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: zorunlu satır eksik: {', '.join(eksik)}. Ticari elektronik iletide gönderen unvanı, MERSİS no, "
                            "iletişim bilgisi ve ret (abonelikten çıkma) imkânı bulunmalıdır.")
    if not re.fullmatch(r"\d{16}", re.sub(r"\s", "", b["mersis"])):
        raise llm.LLMHatasi(f"MERSİS No '{b['mersis']}' 16 haneli olmalı.")
    if not tarih_coz(b["bitis"]):
        raise llm.LLMHatasi(f"Bitiş '{b['bitis']}' tarih olarak okunamadı (GG.AA.YYYY).")
    return b


def bilgi_metni(b: dict) -> str:
    return "\n".join(f"{ALANLAR[a]}: {b[a]}" for a in ("kampanya", "baslangic", "bitis", "indirim", "kod", "hedef", "ton") if b[a]) + \
        ("\nKoşullar:\n" + "\n".join(f"- {k}" for k in b["kosullar"]) if b["kosullar"] else "")


def denetle(t: dict, b: dict, konu_azami: int, konu_sayisi: int) -> list[str]:
    sorun, kaynak = [], bilgi_metni(b)
    if len(t["konu_satirlari"]) < konu_sayisi:
        sorun.append(f"Konu: {konu_sayisi} alternatif istendi, {len(t['konu_satirlari'])} yazıldı")
    for i, k in enumerate(t["konu_satirlari"], 1):
        if len(k) > konu_azami:
            sorun.append(f"Konu {i}: {len(k)} karakter, sınır {konu_azami}")
        buyuk = [w for w in re.findall(r"\b\w{4,}\b", k) if w.isupper() and not w.isdigit() and katla(w) != katla(b["kod"])]
        if buyuk:
            sorun.append(f"Konu {i}: tamamı büyük harf kelime ({', '.join(buyuk)}); spam filtresine takılabilir")
        if k.count("!") > 1:
            sorun.append(f"Konu {i}: birden fazla ünlem")
    alanlar = [("Konu " + str(i), k) for i, k in enumerate(t["konu_satirlari"], 1)] + [("Ön başlık", t["on_baslik"]), ("Başlık", t["baslik"])]
    alanlar += [(f"Gövde {i}", g) for i, g in enumerate(t["govde"], 1)] + [("Buton", t["buton"]), ("Koşul özeti", t["kosul_ozeti"])]
    kk = katla(kaynak)
    izinli = sayilar(kaynak) | {str(x) for d in (tarih_coz(b["baslangic"]), tarih_coz(b["bitis"])) if d for x in (d.day, d.year)}  # "26 Ekim"
    for ad, metin in alanlar:
        fazla = sorted(sayilar(metin) - izinli)
        if fazla:
            sorun.append(f"{ad}: kampanya bilgisinde olmayan sayı / tarih: {', '.join(fazla)}")
        for m in re.finditer(r"%\s*(\d+)", metin):
            if re.search(rf"%\s*{m.group(1)}{EK_YUZDE}", kaynak, re.I) and not re.match(EK_YUZDE, metin[m.end():], re.I):
                sorun.append(f"{ad}: bilgide '%{m.group(1)}'e varan' yazıyor; metinde 'varan / kadar' olmadan '%{m.group(1)}' yanıltıcı")
        for m in IDDIA.finditer(katla(metin)):
            if m.group(0) not in kk:
                sorun.append(f"{ad}: '{m.group(0)}' iddiası kampanya bilgisinde yok")
        if SPAM.search(katla(metin)):
            sorun.append(f"{ad}: spam filtresine takılabilecek ifade ('{SPAM.search(katla(metin)).group(0)}')")
        if URL.search(metin):
            sorun.append(f"{ad}: metne bağlantı yazılmış; bağlantıyı kod butona ekler")
    govde = " ".join(t["govde"] + [t["kosul_ozeti"]])
    bt = tarih_coz(b["bitis"])
    if b["bitis"] not in govde and not re.search(rf"\b{bt.day}\s+{AYLAR[bt.month - 1]}", govde):
        sorun.append(f"Gövde: kampanya bitiş tarihi ({b['bitis']}) yazılmamış")
    if b["kod"] and b["kod"] not in govde:
        sorun.append(f"Gövde: kampanya kodu '{b['kod']}' aynen yazılmamış")
    return sorun


def uret(b: dict, konu_azami: int, konu_sayisi: int, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    mesaj = (f"<kampanya>\n{bilgi_metni(b)}\n</kampanya>\n<konu>{konu_sayisi} alternatif, en fazla {konu_azami} karakter</konu>"
             + duzeltme_blogu(duzeltme, onceki))
    return llm.json_iste(sistem, mesaj, SEMA)


def alt_bilgi(b: dict) -> list[str]:
    return [f"{b['unvan']} · MERSİS No: {b['mersis']}", f"İletişim: {b['iletisim']}",
            "Bu e-postayı ticari elektronik ileti almayı onayladığınız için aldınız. Almak istemiyorsanız aşağıdaki bağlantıdan "
            "dilediğiniz zaman ücretsiz olarak ret bildiriminde bulunabilirsiniz.", f"Abonelikten çık: {b['ret']}"]


def html_yaz(t: dict, b: dict, konu: str) -> str:
    e = html.escape
    govde = "".join(f"<p style=\"margin:0 0 14px\">{e(g)}</p>" for g in t["govde"])
    alt = "<br>".join(e(x) for x in alt_bilgi(b)[:3]) + f"<br><a href=\"{e(b['ret'])}\" style=\"color:#666\">Abonelikten çık</a>"
    return (f"<!doctype html><html lang=\"tr\"><head><meta charset=\"utf-8\"><title>{e(konu)}</title></head>"
            f"<body style=\"margin:0;background:#f4f4f4;font-family:Arial,Helvetica,sans-serif;color:#222\">"
            f"<span style=\"display:none;max-height:0;overflow:hidden\">{e(t['on_baslik'])}</span>"
            f"<table role=\"presentation\" width=\"100%\" cellpadding=\"0\" cellspacing=\"0\"><tr><td align=\"center\" style=\"padding:24px 12px\">"
            f"<table role=\"presentation\" width=\"600\" style=\"max-width:600px;background:#fff\" cellpadding=\"24\"><tr><td>"
            f"<h1 style=\"font-size:24px;margin:0 0 16px\">{e(t['baslik'])}</h1>{govde}"
            f"<p style=\"margin:24px 0\"><a href=\"{e(b['baglanti'])}\" style=\"background:#1d1d1f;color:#fff;padding:12px 22px;text-decoration:none;"
            f"border-radius:6px;display:inline-block\">{e(t['buton'])}</a></p>"
            f"<p style=\"font-size:12px;color:#555\">{e(t['kosul_ozeti'])}</p></td></tr>"
            f"<tr><td style=\"font-size:11px;color:#777;border-top:1px solid #eee\">{alt}</td></tr></table></td></tr></table></body></html>")


def rapor_yaz(cikti: Path, t, b, kalan, tur) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "E-bülten"
    baslik(ws, ["Bölüm", "Metin", "Karakter", "Seçim"], (14, 90, 9, 8))
    for i, k in enumerate(t["konu_satirlari"], 1):
        ws.append([f"Konu {i}", k, len(k), ""])
        boya(ws, [2], AI_DOLGU)
        boya(ws, [4], SARI)
    for ad, m in [("Ön başlık", t["on_baslik"]), ("Başlık", t["baslik"])] + [(f"Gövde {i}", g) for i, g in enumerate(t["govde"], 1)] + \
            [("Buton", t["buton"]), ("Koşul özeti", t["kosul_ozeti"])]:
        ws.append([ad, m, len(m), ""])
        boya(ws, [2], AI_DOLGU)
    for x in alt_bilgi(b):
        ws.append(["Alt bilgi (kod)", x, "", ""])
    son_not(ws, "Önemli: Göndermeden önce test e-postası alın; alıcıların İYS'de onaylı olduğunu ve kampanya koşullarının güncel olduğunu "
                "kontrol edin.", kalan, tur)
    hizala(ws)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    konu = t["konu_satirlari"][0] if t["konu_satirlari"] else b["kampanya"]
    cikti.with_suffix(".html").write_text(html_yaz(t, b, konu), encoding="utf-8")
    duz = [f"Konu: {konu}", f"Ön başlık: {t['on_baslik']}", "", t["baslik"], ""] + [g + "\n" for g in t["govde"]]
    duz += [f"{t['buton']}: {b['baglanti']}", "", t["kosul_ozeti"], "", "—"] + alt_bilgi(b)
    cikti.with_suffix(".txt").write_text("\n".join(duz), encoding="utf-8")


def calistir(girdi: Path, cikti: Path, konu_azami: int = 60, konu_sayisi: int = 5, tur: int = 2, evet: bool = False) -> dict:
    b = bilgiyi_oku(girdi)
    print(f"[OK] {b['kampanya']} · bitiş {b['bitis']}{' · kod ' + b['kod'] if b['kod'] else ''} · {len(b['kosullar'])} koşul")
    llm.onay_al("Kampanya bilgisi ve koşullar gönderilecek (gönderen unvanı, MERSİS no ve iletişim bilgileri gönderilmez).", evet)
    t, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(b, konu_azami, konu_sayisi, d, o), lambda x: denetle(x, b, konu_azami, konu_sayisi), tur)
    rapor_yaz(cikti, t, b, kalan, yapilan)
    return {"bilgi": b, "taslak": t, "kalan": kalan, "tur": yapilan}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    p = argparse.ArgumentParser(description="Kampanya bilgisinden konu satırı alternatifleriyle e-bülten metni hazırlar (göndermez).")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "kampanya.txt",
                   help="Kampanya bilgisi (.txt/.docx): Gönderen Unvan, MERSİS No, İletişim, Ret Bağlantısı, Kampanya, Başlangıç, Bitiş, İndirim, Kod, "
                        "Bağlantı satırları ve 'Koşullar' bölümü")
    p.add_argument("--konu-azami", type=int, default=60, help="Konu satırı için en fazla karakter (varsayılan 60)")
    p.add_argument("--konu-sayisi", type=int, default=5, help="Konu satırı alternatifi sayısı (varsayılan 5)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "e_bulten.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.cikti, a.konu_azami, a.konu_sayisi, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    for k in s["taslak"]["konu_satirlari"]:
        print(f"[OK] Konu: {k}")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Taslak: {a.cikti.with_suffix('.html').resolve()} (+ .txt, .xlsx)")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
