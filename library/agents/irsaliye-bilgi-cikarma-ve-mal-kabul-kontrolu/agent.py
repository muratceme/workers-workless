"""
İrsaliye Bilgi Çıkarma ve Mal Kabul Kontrolü — Gerisi Bende AI Agent
Depo › Depo Sorumlusu (hastane eczanesi ve lojistik depo rolleri de kullanır)

PDF / taranıp metne çevrilmiş / e-postayla gelen sevk irsaliyelerinden mal kabul kaydı çıkarır ve açık siparişlerle
karşılaştırır.
1. Kod irsaliye belgelerini ve (varsa) açık sipariş listesini okur; telefon, e-posta, IBAN ve kimlik numaralarını maskeler.
2. Model her irsaliyenin başlığını (no, tarih, sevk eden, sipariş referansı) ve satırlarını (stok kodu, ürün adı, miktar,
   birim, lot / seri, son kullanma tarihi) belgedeki yazımla çıkarır ve her satırı açık sipariş listesindeki malzemeyle
   eşleştirir (tedarikçi ürünü kendi adı ve koduyla yazar).
3. Kod çıkarımı denetler: metinler, lot ve SKT belgede aynen geçiyor mu; miktar belgede yazıyor mu; eşleşen malzeme kodu
   o siparişte var mı; aynı sipariş kalemi iki satıra verilmiş mi. Sorun varsa modele geri gönderilir (en fazla --tur kez).
4. Kod mal kabul bulgularını hesaplar: siparişte olmayan ürün, fazla teslim, kısmi teslim (kalan miktar), son kullanma
   tarihi geçmiş ya da --skt-ay içinde dolacak ürün, SKT'li üründe lot yok, irsaliyede sipariş referansı yok.
Çıktının ilk sayfası "Son Kullanma Tarihi Takibi" kod bloğunun girdi sütunlarını içerir.

Kullanım:
    python agent.py                                                           # örnek 2 irsaliye
    python agent.py --girdi ./irsaliyeler --siparis acik_siparisler.xlsx --skt-ay 6
"""
from __future__ import annotations

import argparse
import calendar
import csv
import json
import re
import sys
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
ORNEK_TARIH = date(2026, 10, 16)
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
RENK = {"Hata": "FDE2E1", "Dikkat": "FFF4CE", "Bilgi": "E8F0FE", "Red / İnceleme": "FDE2E1", "Kontrol": "FFF4CE", "Kabul": "E3F5E1"}
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1 or (s.count(".") == 1 and len(s.split(".")[1]) == 3):
        s = s.replace(".", "")
    try:
        return Decimal(s) if s not in ("", "-", ".") else None
    except InvalidOperation:
        return None


def sayilar(metin: str) -> set[Decimal]:
    return {v for x in re.findall(r"\d+(?:[.,]\d+)*", str(metin)) if (v := para(x)) is not None}


def skt_coz(x) -> date | None:
    """'09/2028', '09.2028', '2028-09' → ayın son günü; '30.06.2027', '2027-06-30' → o gün."""
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    m = re.fullmatch(r"(\d{1,2})[./-](\d{1,2})[./-](\d{4})", s)
    if m:
        try:
            return date(int(m[3]), int(m[2]), int(m[1]))
        except ValueError:
            return None
    m = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return date(int(m[1]), int(m[2]), int(m[3]))
    m = re.fullmatch(r"(\d{1,2})[./](\d{4})", s) or re.fullmatch(r"(\d{4})-(\d{1,2})", s)
    if m:
        ay, yil = (int(m[1]), int(m[2])) if len(m[1]) <= 2 else (int(m[2]), int(m[1]))
        if 1 <= ay <= 12:
            return date(yil, ay, calendar.monthrange(yil, ay)[1])
    return None


def tarih_coz(x) -> date | None:
    d = skt_coz(x)
    return d if d and re.search(r"\d{1,2}[./-]\d{1,2}[./-]\d{4}|\d{4}-\d{2}-\d{2}", str(x)) else None


def ay_ekle(d: date, n: int) -> date:
    y, a = divmod(d.month - 1 + n, 12)
    yil, ay = d.year + y, a + 1
    return date(yil, ay, min(d.day, calendar.monthrange(yil, ay)[1]))


def tr(x: Decimal) -> str:
    s = f"{x:,.3f}".rstrip("0").rstrip(".")
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


# ----------------------------------------------------------------------------
# Girdi
# ----------------------------------------------------------------------------

SIPARIS_SUTUN = {"no": ("siparis no", "siparis numarasi", "po no", "satin alma siparisi"), "kod": ("malzeme kodu", "stok kodu", "urun kodu", "ilac kodu", "kod"),
                 "ad": ("malzeme adi", "urun adi", "malzeme", "urun", "ilac adi"), "miktar": ("siparis miktari", "miktar"),
                 "teslim": ("teslim alinan", "onceki teslim", "gelen miktar", "teslim alinan miktar"), "birim": ("birim",)}


def siparis_oku(yol: Path | None) -> list[dict]:
    if yol is None:
        return []
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: sipariş listesi bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, read_only=True, data_only=True)
        s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=metin.split("\n")[0].count)))
    s = [r for r in s if any(c not in (None, "") for c in r)]
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: sipariş listesi boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in es), None) for a, es in SIPARIS_SUTUN.items()}
    eksik = [SIPARIS_SUTUN[a][0] for a in ("no", "kod", "ad", "miktar") if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun yok: {', '.join(eksik)}. Başlıklar: {s[0]}")
    al = lambda r, a: "" if idx[a] is None or idx[a] >= len(r) or r[idx[a]] is None else str(r[idx[a]]).strip()  # noqa: E731
    out = []
    for r in s[1:]:
        m, t = para(al(r, "miktar")), para(al(r, "teslim")) or Decimal(0)
        out.append({"no": al(r, "no"), "kod": al(r, "kod"), "ad": al(r, "ad"), "miktar": m, "onceki": t, "birim": al(r, "birim"),
                    "kalan": None if m is None else m - t})
    return [x for x in out if x["kod"]]


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

SATIR = {"type": "object", "properties": {"stok_kodu": S, "urun_adi": S, "miktar": S, "birim": S, "lot": S, "skt": S, "siparis_kodu": S},
         "required": ["stok_kodu", "urun_adi", "miktar", "birim", "lot", "skt", "siparis_kodu"], "additionalProperties": False}
SEMA = {"type": "object", "properties": {
    "irsaliye_var": {"type": "boolean"}, "irsaliye_no": S, "tarih": S, "sevk_eden": S, "alici": S, "siparis_no": S,
    "satirlar": {"type": "array", "items": SATIR}, "notlar": {"type": "array", "items": S}},
    "required": ["irsaliye_var", "irsaliye_no", "tarih", "sevk_eden", "alici", "siparis_no", "satirlar", "notlar"], "additionalProperties": False}


def denetle(c: dict, metin: str, siparis: list[dict]) -> list[str]:
    if not c["irsaliye_var"]:
        return []
    sorun, say = [], sayilar(metin)
    for a, ad in (("irsaliye_no", "İrsaliye no"), ("tarih", "Tarih"), ("sevk_eden", "Sevk eden"), ("alici", "Alıcı"), ("siparis_no", "Sipariş no")):
        if c[a].strip() and not metinde_var(c[a], metin):
            sorun.append(f"İrsaliye: {ad} '{c[a][:50]}' belgede geçmiyor (belgedeki yazımı aynen kullanın)")
    if c["tarih"].strip() and tarih_coz(c["tarih"]) is None:
        sorun.append(f"İrsaliye: tarih '{c['tarih']}' GG.AA.YYYY biçiminde değil")
    if not c["satirlar"]:
        sorun.append("İrsaliye: satır yok; belgedeki ürün satırlarını yazın")
    sip_kod = {x["kod"] for x in siparis}
    bu_sip = {x["kod"] for x in siparis if x["no"] == c["siparis_no"].strip()}
    gorulen: dict[str, int] = {}
    for j, s in enumerate(c["satirlar"], 1):
        on = f"Satır {j}"
        if not metinde_var(s["urun_adi"], metin):
            sorun.append(f"{on}: ürün adı '{s['urun_adi'][:50]}' belgede geçmiyor")
        for a, ad in (("stok_kodu", "Stok kodu"), ("lot", "Lot"), ("skt", "SKT")):
            if s[a].strip() and not metinde_var(s[a], metin):
                sorun.append(f"{on}: {ad} '{s[a]}' belgede geçmiyor")
        if skt_coz(s["lot"].strip()):
            sorun.append(f"{on}: lot '{s['lot']}' bir tarih gibi görünüyor; sütun kaymış olabilir, SKT alanına mı ait?")
        if s["skt"].strip() and skt_coz(s["skt"]) is None:
            sorun.append(f"{on}: SKT '{s['skt']}' tarih olarak okunamadı (belgedeki biçimi koruyun: AA/YYYY ya da GG.AA.YYYY)")
        m = para(s["miktar"])
        if m is None:
            sorun.append(f"{on}: miktar '{s['miktar']}' okunamadı")
        elif m not in say:
            sorun.append(f"{on}: miktar {s['miktar']} belgede yazmıyor")
        k = s["siparis_kodu"].strip()
        if not k:
            continue
        if k not in sip_kod:
            sorun.append(f"{on}: sipariş kodu '{k}' sipariş listesinde yok; eşleşme yoksa boş bırakın")
        elif bu_sip and k not in bu_sip:
            sorun.append(f"{on}: {k} irsaliyedeki sipariş ({c['siparis_no']}) altında değil")
        if k in gorulen:          # aynı ürün farklı lotlarla iki satırda gelebilir; farklı ürün, lotsuz ya da aynı lotla gelemez
            o = c["satirlar"][gorulen[k] - 1]
            if not (katla(s["urun_adi"]) == katla(o["urun_adi"]) and s["lot"].strip() and o["lot"].strip() and katla(s["lot"]) != katla(o["lot"])):
                sorun.append(f"{on}: sipariş kodu {k} satır {gorulen[k]} ile aynı; her sipariş kalemi bir satıra eşleşir (aynı ürün farklı lotlarla "
                             "geldiyse iki satırın lotunu yazın)")
        gorulen.setdefault(k, j)
    return sorun


def uret(metin: str, siparis: list[dict], bugun: date, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    liste = "\n".join(f"{x['no']} | {x['kod']} | {x['ad']} | kalan {tr(x['kalan']) if x['kalan'] is not None else '?'} {x['birim']}".rstrip() for x in siparis)
    mesaj = (f"<acik_siparisler>\n{liste or '(verilmedi: siparis_kodu alanlarını boş bırakın)'}\n</acik_siparisler>\n<bugun>{bugun:%d.%m.%Y}</bugun>\n"
             f"<irsaliye>\n{metin}\n</irsaliye>")
    if duzeltme:
        mesaj += ("\n<duzeltme>\nÖnceki çıkarım: " + json.dumps(onceki, ensure_ascii=False) + "\nKodun bulduğu sorunlar:\n"
                  + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")
    return llm.json_iste(sistem, mesaj, SEMA)


def duzeltme_dongusu(uret_f, denetle_f, tur: int) -> tuple[dict, list[str], int]:
    taslak = uret_f(None, None)
    kalan = denetle_f(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret_f(kalan, taslak)
        kalan = denetle_f(taslak)
    return taslak, kalan, yapilan


def mal_kabul(irsaliyeler: list[dict], siparis: list[dict], bugun: date, skt_ay: int) -> None:
    """Kod: her satır için bulgular ve kabul durumu (yerinde ekler). Aynı kaleme birden çok irsaliye / lot gelirse toplanır."""
    sip = {(x["no"], x["kod"]): x for x in siparis}
    kod_sip = defaultdict(list)
    for x in siparis:
        kod_sip[x["kod"]].append(x)
    toplam: dict[tuple, Decimal] = defaultdict(Decimal)
    sinir = ay_ekle(bugun, skt_ay)
    for d in irsaliyeler:
        c = d["c"]
        if not c["irsaliye_var"]:
            continue
        d["bulgular"] = []
        if siparis and not c["siparis_no"].strip():
            d["bulgular"].append(("Dikkat", "İrsaliyede sipariş referansı yok; sipariş eşleşmesini kontrol edin"))
        for s in c["satirlar"]:
            b = []
            k, m = s["siparis_kodu"].strip(), para(s["miktar"]) or Decimal(0)
            kalem = sip.get((c["siparis_no"].strip(), k)) or (kod_sip[k][0] if len(kod_sip.get(k, [])) == 1 else None)
            if siparis and not k:
                b.append(("Hata", "Açık siparişlerde yok: kabul etmeden önce satın almaya sorun"))
            if kalem:
                anahtar = (kalem["no"], kalem["kod"])
                toplam[anahtar] += m
                s["_kalem"] = kalem
            skt = skt_coz(s["skt"])
            if skt and skt < bugun:
                b.append(("Hata", f"Son kullanma tarihi geçmiş ({skt:%d.%m.%Y})"))
            elif skt and skt <= sinir:
                b.append(("Dikkat", f"Son kullanma tarihi {skt_ay} ay içinde doluyor ({skt:%d.%m.%Y}); kabul kriterinizi kontrol edin"))
            if s["skt"].strip() and not s["lot"].strip():
                b.append(("Dikkat", "Son kullanma tarihli üründe lot / seri no yok (izlenebilirlik)"))
            s["_bulgu"] = b
    for d in irsaliyeler:
        if not d["c"]["irsaliye_var"]:
            continue
        for s in d["c"]["satirlar"]:
            kalem = s.get("_kalem")
            if kalem and kalem["kalan"] is not None:
                gelen = toplam[(kalem["no"], kalem["kod"])]
                fark = gelen - kalem["kalan"]
                s["_fark"] = fark
                if fark > 0:
                    s["_bulgu"].append(("Dikkat", f"Fazla teslim: kalan sipariş {tr(kalem['kalan'])}, gelen {tr(gelen)} (+{tr(fark)})"))
                elif fark < 0:
                    s["_bulgu"].append(("Bilgi", f"Kısmi teslim: kalan sipariş {tr(kalem['kalan'])}, gelen {tr(gelen)}; açık kalan {tr(-fark)}"))
            onem = {o for o, _ in s["_bulgu"]}
            s["_durum"] = "Red / İnceleme" if "Hata" in onem else "Kontrol" if "Dikkat" in onem else "Kabul"


# ----------------------------------------------------------------------------
# Çıktı
# ----------------------------------------------------------------------------

def _baslik(ws, basliklar, genislik) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font, h.alignment = BASLIK_DOLGU, BASLIK_YAZI, UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def rapor_yaz(cikti: Path, irsaliyeler: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Teslim Alınanlar"
    _baslik(ws, ["Malzeme Kodu", "Malzeme Adı", "Lot No", "Son Kullanma Tarihi", "Miktar", "Birim", "Sipariş No", "Sipariş Kalanı", "Fark",
                 "İrsaliye No", "Tedarikçi Kodu", "İrsaliyedeki Ad", "Durum", "Bulgular", "Kaynak Dosya"],
            (12, 36, 12, 12, 9, 7, 14, 10, 9, 20, 12, 40, 14, 60, 28))
    for d in irsaliyeler:
        c = d["c"]
        if not c["irsaliye_var"]:
            continue
        for s in c["satirlar"]:
            kalem = s.get("_kalem")
            skt = skt_coz(s["skt"])
            ws.append([kalem["kod"] if kalem else "", kalem["ad"] if kalem else s["urun_adi"], s["lot"], skt, _f(s["miktar"]), s["birim"],
                       kalem["no"] if kalem else c["siparis_no"], _f(kalem["kalan"]) if kalem else None, _f(s.get("_fark")), c["irsaliye_no"],
                       s["stok_kodu"], s["urun_adi"], s["_durum"], "\n".join(f"{o}: {m}" for o, m in s["_bulgu"]), d["dosya"]])
            n = ws.max_row
            ws.cell(n, 4).number_format = "DD.MM.YYYY"
            for j in (1, 3, 4, 5, 6, 11, 12):
                ws.cell(n, j).fill = AI_DOLGU
            ws.cell(n, 13).fill = PatternFill("solid", fgColor=RENK[s["_durum"]])
            for h in ws[n]:
                h.alignment = UST
    i = wb.create_sheet("İrsaliyeler")
    _baslik(i, ["Kaynak Dosya", "İrsaliye No", "Tarih", "Sevk Eden", "Alıcı", "Sipariş No", "Satır", "Kabul", "Kontrol", "Red / İnceleme", "Bulgular",
                "Model Notları", "Düzeltme Turu", "Son Kontrol"], (28, 20, 11, 30, 30, 14, 6, 6, 7, 8, 40, 40, 8, 40))
    for d in irsaliyeler:
        c = d["c"]
        if not c["irsaliye_var"]:
            i.append([d["dosya"], "", "", "", "", "", 0, 0, 0, 0, "Belgede sevk irsaliyesi bulunamadı", "\n".join(c["notlar"]), d["tur"], ""])
            continue
        dur = [s["_durum"] for s in c["satirlar"]]
        i.append([d["dosya"], c["irsaliye_no"], tarih_coz(c["tarih"]) or c["tarih"], c["sevk_eden"], c["alici"], c["siparis_no"], len(dur),
                  dur.count("Kabul"), dur.count("Kontrol"), dur.count("Red / İnceleme"), "\n".join(f"{o}: {m}" for o, m in d["bulgular"]),
                  "\n".join(c["notlar"]), d["tur"], "; ".join(d["kalan"]) or "Geçti"])
        i.cell(i.max_row, 3).number_format = "DD.MM.YYYY"
        i.cell(i.max_row, 14).fill = KIRMIZI if d["kalan"] else YESIL
        for h in i[i.max_row]:
            h.alignment = UST
    k = wb.create_sheet("Bilgi")
    k.append([f"Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    k.append(["Önemli: Mor hücreler yapay zekâ ile çıkarılmıştır. Malı fiziksel sayım ve irsaliye aslıyla karşılaştırmadan kabul etmeyin. "
              "Fark, kısmi / fazla teslim ve son kullanma kontrolleri kod tarafından hesaplanmıştır. 'Teslim Alınanlar' sayfası "
              "'Son Kullanma Tarihi Takibi' kod bloğunun girdi sütunlarını içerir."])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def _f(x):
    v = para(x) if not isinstance(x, Decimal) else x
    return None if v is None else float(v)


def dosyalar(girdi: Path) -> list[Path]:
    if girdi.is_dir():
        l = sorted(p for p in girdi.iterdir() if p.is_file() and p.suffix.lower() in belge.DESTEKLENEN)
        if not l:
            raise belge.BelgeHatasi(f"{girdi}: klasörde .pdf, .docx veya .txt dosyası yok.")
        return l
    if not girdi.exists():
        raise belge.BelgeHatasi(f"{girdi}: dosya bulunamadı.")
    return [girdi]


def calistir(girdi: Path, cikti: Path, siparis_yolu: Path | None = None, bugun: date = ORNEK_TARIH, skt_ay: int = 6, tur: int = 2,
             evet: bool = False) -> dict:
    siparis = siparis_oku(siparis_yolu)
    l = dosyalar(girdi)
    print(f"[OK] {len(l)} irsaliye belgesi · açık sipariş kalemi: {len(siparis)}")
    llm.onay_al(f"{len(l)} irsaliyenin metni ve açık sipariş listesi gönderilecek (telefon, e-posta, IBAN, TCKN maskeli).", evet)
    irsaliyeler = []
    for y in l:
        ham = belge.metin_oku(y)
        if len(ham.strip()) < 20:
            raise belge.BelgeHatasi(f"{y.name}: metin bulunamadı.")
        metin = llm.maskele(ham)
        c, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(metin, siparis, bugun, d, o), lambda x: denetle(x, metin, siparis), tur)
        irsaliyeler.append({"dosya": y.name, "c": c, "kalan": kalan, "tur": yapilan, "bulgular": []})
    mal_kabul(irsaliyeler, siparis, bugun, skt_ay)
    rapor_yaz(cikti, irsaliyeler)
    return {"irsaliyeler": irsaliyeler}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Sevk irsaliyelerinden mal kabul kaydı çıkarır ve açık siparişlerle karşılaştırır.")
    p.add_argument("--girdi", type=Path, default=ORNEK / "irsaliyeler", help="İrsaliye belgesi ya da klasörü (.pdf, .docx, .txt)")
    p.add_argument("--siparis", type=Path, default=None, help="Açık siparişler (.xlsx/.csv: Sipariş No, Malzeme Kodu, Malzeme Adı, Sipariş Miktarı, Teslim Alınan)")
    p.add_argument("--tarih", default=None, help="Kabul tarihi GG.AA.YYYY (varsayılan bugün; örnekte 16.10.2026)")
    p.add_argument("--skt-ay", type=int, default=6, help="Son kullanma tarihi bu kadar ay içinde dolacaksa uyar (varsayılan 6)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "mal_kabul.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    ornek = a.girdi == ORNEK / "irsaliyeler"
    siparis = a.siparis or (ORNEK / "acik_siparisler.csv" if ornek else None)
    bugun = tarih_coz(a.tarih) if a.tarih else (ORNEK_TARIH if ornek else date.today())
    if bugun is None or a.skt_ay < 0:
        print("[X] --tarih GG.AA.YYYY biçiminde, --skt-ay 0 ya da büyük olmalı")
        return 2
    try:
        s = calistir(a.girdi, a.cikti, siparis, bugun, a.skt_ay, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    for d in s["irsaliyeler"]:
        c = d["c"]
        if not c["irsaliye_var"]:
            print(f"[!] {d['dosya']}: sevk irsaliyesi bulunamadı")
            continue
        dur = [x["_durum"] for x in c["satirlar"]]
        print(f"[OK] {c['irsaliye_no']} ({c['sevk_eden']}): {len(dur)} satır · kabul {dur.count('Kabul')}, kontrol {dur.count('Kontrol')}, "
              f"red / inceleme {dur.count('Red / İnceleme')}")
        for x in c["satirlar"]:
            for o, m in x["_bulgu"]:
                if o == "Hata":
                    print(f"[!]   {x['urun_adi']}: {m}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
