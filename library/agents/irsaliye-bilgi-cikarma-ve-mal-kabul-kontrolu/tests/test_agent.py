"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import json
import os
import sys
import tempfile
import unittest
from datetime import date, datetime
from decimal import Decimal as D
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
IRS = ORNEK / "irsaliyeler"
SIPARIS = ORNEK / "acik_siparisler.csv"
BUGUN = date(2026, 10, 16)
CEVAP = json.loads((Path(__file__).parent / "sahte_cevaplar.json").read_text(encoding="utf-8"))
MESAJLAR = []


def ecza_hatali() -> dict:
    """Sık model hataları: lot sütunu kaydığı için SKT'yi lot olarak okumak, miktarı siparişten almak, olmayan kod uydurmak."""
    c = copy.deepcopy(CEVAP["ecza"])
    c["satirlar"][2].update(lot="05/2028", skt="")
    c["satirlar"][0]["miktar"] = "600"
    c["satirlar"][3]["siparis_kodu"] = "IL-0999"
    c["satirlar"][1]["siparis_kodu"] = "IL-0101"
    return c


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    if "ÖRNEK AMBALAJ" in kullanici:
        return copy.deepcopy(CEVAP["ambalaj"])
    return copy.deepcopy(CEVAP["ecza"]) if "<duzeltme>" in kullanici else ecza_hatali()


class KodTesti(unittest.TestCase):
    def test_tarih_ve_siparis(self):
        self.assertEqual([agent.skt_coz(x) for x in ("09/2028", "02.2027", "2027-06", "30.06.2027", "2027-06-30", "13/2027", "yakında")],
                         [date(2028, 9, 30), date(2027, 2, 28), date(2027, 6, 30), date(2027, 6, 30), date(2027, 6, 30), None, None])
        self.assertEqual((agent.tarih_coz("15.10.2026"), agent.tarih_coz("10/2026")), (date(2026, 10, 15), None))
        self.assertEqual((agent.ay_ekle(date(2026, 8, 31), 6), agent.ay_ekle(date(2026, 10, 16), 6)), (date(2027, 2, 28), date(2027, 4, 16)))
        s = agent.siparis_oku(SIPARIS)
        self.assertEqual([(x["kod"], x["kalan"]) for x in s][:3], [("IL-0101", D(600)), ("IL-0233", D(200)), ("IL-0410", D(400))])
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "s.csv"
            y.write_text("Kod;Miktar\nA;1\n", encoding="utf-8")
            with self.assertRaises(llm.LLMHatasi):
                agent.siparis_oku(y)

    def test_denetim(self):
        s = agent.siparis_oku(SIPARIS)
        m = llm.maskele((IRS / "kurgu_ecza_deposu_irsaliye.txt").read_text(encoding="utf-8"))
        self.assertEqual(agent.denetle(CEVAP["ecza"], m, s), [])
        h = agent.denetle(ecza_hatali(), m, s)
        self.assertEqual(len(h), 4)
        for parca in ("Satır 1: miktar 600 belgede yazmıyor", "Satır 2: sipariş kodu IL-0101 satır 1 ile aynı", "Satır 3: lot '05/2028' bir tarih", "IL-0999"):
            self.assertTrue(any(parca in x for x in h), parca)
        self.assertFalse(any("Satır 3: Lot" in x for x in h))                 # '05/2028' belgede geçiyor; hata SKT boşluğunda değil
        c = copy.deepcopy(CEVAP["ecza"])                                      # aynı ürün aynı lotla iki satırda: hata
        c["satirlar"][1] = c["satirlar"][0] | {"miktar": "200"}
        self.assertTrue(any("satır 1 ile aynı" in x for x in agent.denetle(c, m, s)))
        c["satirlar"][1]["lot"] = "SFT2508"                                   # aynı ürün farklı lotla: hata değil
        self.assertFalse(any("satır 1 ile aynı" in x for x in agent.denetle(c, m, s)))
        a = llm.maskele((IRS / "ornek_ambalaj_irsaliye.txt").read_text(encoding="utf-8"))
        c = copy.deepcopy(CEVAP["ambalaj"])
        c["satirlar"][0]["siparis_kodu"] = "IL-0101"
        self.assertIn("SA-2026-0698", agent.denetle(c, a, s)[0])              # kod başka siparişin kalemi
        self.assertEqual(agent.denetle({"irsaliye_var": False}, a, s), [])

    def test_mal_kabul(self):
        s = agent.siparis_oku(SIPARIS)
        irs = [{"dosya": "a", "c": copy.deepcopy(CEVAP["ecza"]), "bulgular": []}, {"dosya": "b", "c": copy.deepcopy(CEVAP["ambalaj"]), "bulgular": []}]
        agent.mal_kabul(irs, s, BUGUN, 6)
        sat = [x for d in irs for x in d["c"]["satirlar"]]
        self.assertEqual([x["_durum"] for x in sat], ["Kontrol", "Kontrol", "Kontrol", "Red / İnceleme", "Kabul", "Kabul", "Kabul"])
        self.assertEqual([x.get("_fark") for x in sat], [D(50), D(0), D(0), None, D(-800), D(0), D(0)])
        self.assertIn("6 ay içinde", sat[1]["_bulgu"][0][1])
        agent.mal_kabul(irs, s, date(2027, 4, 1), 0)                           # SKT geçmiş
        self.assertEqual(sat[1]["_bulgu"][0], ("Hata", "Son kullanma tarihi geçmiş (31.03.2027)"))
        iki = [{"dosya": "a", "c": copy.deepcopy(CEVAP["ambalaj"]), "bulgular": []}, {"dosya": "b", "c": copy.deepcopy(CEVAP["ambalaj"]), "bulgular": []}]
        agent.mal_kabul(iki, s, BUGUN, 6)                                      # iki irsaliye aynı kaleme: toplam 2.400 > 2.000
        self.assertEqual(iki[1]["c"]["satirlar"][0]["_bulgu"][0], ("Dikkat", "Fazla teslim: kalan sipariş 2.000, gelen 2.400 (+400)"))
        bos = [{"dosya": "a", "c": copy.deepcopy(CEVAP["ambalaj"]) | {"siparis_no": ""}, "bulgular": []}]
        agent.mal_kabul(bos, s, BUGUN, 6)
        self.assertEqual(bos[0]["bulgular"][0][0], "Dikkat")


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ornek_ve_duzeltme(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "m.xlsx"
            self.assertEqual(agent.main(["--evet", "--cikti", str(cikti)]), 0)
            self.assertEqual(len(MESAJLAR), 3)                           # 2 irsaliye + ecza için 1 düzeltme
            self.assertIn("SA-2026-0712 | IL-0233 | Seftriakson 1 g IM/IV flakon | kalan 200 flakon", MESAJLAR[0])
            self.assertIn("miktar 600 belgede yazmıyor", MESAJLAR[1].split("<duzeltme>")[1])
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Teslim Alınanlar", "İrsaliyeler", "Bilgi"])
            ws = wb["Teslim Alınanlar"]
            self.assertEqual([c.value for c in ws[1]][:5], ["Malzeme Kodu", "Malzeme Adı", "Lot No", "Son Kullanma Tarihi", "Miktar"])
            self.assertEqual([ws.cell(2, j).value for j in (1, 2, 3, 4, 5, 9, 13)],
                             ["IL-0101", "Parasetamol 500 mg tablet (20'li kutu)", "PRC2510", datetime(2028, 9, 30), 650, 50, "Kontrol"])
            self.assertEqual([ws.cell(5, j).value for j in (1, 2, 13)], [None, "ENOKSAPARİN 4000 IU ŞIRINGA", "Red / İnceleme"])
            self.assertEqual([r[7:10] for r in wb["İrsaliyeler"].iter_rows(min_row=2, values_only=True)], [(0, 3, 1), (3, 0, 0)])

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_siparissiz_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(IRS / "ornek_ambalaj_irsaliye.txt", Path(tmp) / "m.xlsx", None, BUGUN, evet=True)
            self.assertEqual([x["_durum"] for x in s["irsaliyeler"][0]["c"]["satirlar"]], ["Kabul"] * 3)
            self.assertEqual(agent.main(["--tarih", "yarın", "--evet"]), 2)
            self.assertEqual(agent.main(["--skt-ay", "-1", "--evet"]), 2)
            self.assertEqual(agent.main(["--girdi", str(Path(tmp) / "yok.pdf"), "--evet"]), 1)
            self.assertEqual(agent.main(["--siparis", str(Path(tmp) / "yok.csv"), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
