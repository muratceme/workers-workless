"""
Teminat Mektubu Metin Kontrolü — Workers / Workless AI Agent
Bankacılık Operasyonları › Dış Ticaret Operasyon Uzmanı

1. Kod talep formunu (lehtar, borçlu, tutar, para birimi, vade, sözleşme), bankanın standart metnini ({LEHTAR}, {TUTAR} …
   yer tutuculu) ve lehtarın / müşterinin istediği metni okur. Standart metin formdaki bilgilerle doldurulur; iki metin
   cümlelere bölünür (S1, S2 … ve T1, T2 …). Lehtar ve borçlu unvanları modele gönderilmez ({LEHTAR}, {BORCLU}).
2. Kod alanları karşılaştırır (model yok):
   - Talep metnindeki lehtar, borçlu, sözleşme ve tutar formla aynı mı; başka tutar ya da para birimi var mı
   - Yazıyla tutar rakamla aynı mı (ör. 2.450.000,00 → "İki Milyon Dört Yüz Elli Bin Türk Lirası")
   - Geçerlilik tarihi formdaki vadeyle aynı mı; vade düzenleme tarihinden sonra mı
   - Standart metindeki hangi cümleler talep metninde yok
3. Model talep metnindeki her cümleyi standart cümleyle eşleştirir: Aynı / İfade farkı / Esaslı sapma / Standartta yok,
   risk düzeyi ve açıklama.
4. Kod denetler:
   - Her talep cümlesi bir kez sınıflandırılmış mı; eşleştirilen standart cümle var mı
   - "Aynı" denen cümle standart cümleyle gerçekten aynı mı (benzerlik)
   - Faiz, kendiliğinden uzama, devir / temlik, yabancı mahkeme / tahkim, süresizlik, kısmi ödeme gibi riskli hükümler
     (standartta yoksa) esaslı sapma ve en az orta risk olarak işaretlenmiş mi
   - Açıklamada metinlerde olmayan sayı var mı
   Sorun varsa modele geri gönderilir (en fazla --tur kez).

Kullanım:
    python agent.py                                                     # örnek (kurgusal) kesin teminat talebi
    python agent.py --form talep.txt --standart banka.txt --metin lehtar_metni.docx
"""
from __future__ import annotations

import argparse
import csv
import difflib
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
FORM = {"lehtar": "Lehtar", "borclu": "Borçlu", "tur": "Tür", "tutar": "Tutar", "para": "Para Birimi", "vade": "Vade",
        "duzenleme": "Düzenleme Tarihi", "sozlesme": "Sözleşme"}
PARA_ADI = {"tl": "Türk Lirası", "try": "Türk Lirası", "usd": "ABD Doları", "eur": "Avro", "euro": "Avro", "gbp": "İngiliz Sterlini"}
PARA_KELIME = ("turk", "lira", "tl", "abd", "dolar", "amerikan", "avro", "euro", "ingiliz", "sterlin")
DURUMLAR = ["Aynı", "İfade farkı", "Esaslı sapma", "Standartta yok"]
RISKLER = ["Düşük", "Orta", "Yüksek"]
RISKLI = [("faiz", r"\bfaiz"), ("kendiliğinden uzama", r"kendiliginden|otomatik\w* (olarak )?uza|uzamis say"), ("devir / temlik", r"\bdevir|temlik|devred"),
          ("yabancı hukuk / mahkeme / tahkim", r"yabanci|londra|ingiliz hukuk|isvicre|tahkim|\bicc\b"), ("süresiz", r"suresiz"),
          ("kısmi / tekrarlı ödeme", r"kismen|kisim kisim|birden fazla talep"), ("ödeme süresi / ek koşul", r"\d+ (is )?gun(u)? icinde|sartiyla|kosuluyla")]
BIRLER = ["", "Bir", "İki", "Üç", "Dört", "Beş", "Altı", "Yedi", "Sekiz", "Dokuz"]
ONLAR = ["", "On", "Yirmi", "Otuz", "Kırk", "Elli", "Altmış", "Yetmiş", "Seksen", "Doksan"]
GRUPLAR = ["", "Bin", "Milyon", "Milyar", "Trilyon"]

SEMA = {
    "type": "object",
    "properties": {
        "eslesmeler": {"type": "array", "items": {"type": "object", "properties": {
            "talep": S, "standart": S, "durum": {"type": "string", "enum": DURUMLAR}, "risk": {"type": "string", "enum": RISKLER}, "aciklama": S},
            "required": ["talep", "standart", "durum", "risk", "aciklama"], "additionalProperties": False}},
        "genel_degerlendirme": S,
    },
    "required": ["eslesmeler", "genel_degerlendirme"],
    "additionalProperties": False,
}


def yaziyla(n: int) -> str:
    """2450000 → 'İki Milyon Dört Yüz Elli Bin' (Türkçe bankacılık yazımı; 'Bir Bin' değil 'Bin')."""
    if n == 0:
        return "Sıfır"
    parca, i = [], 0
    while n:
        g = n % 1000
        if g:
            uc = ([("" if g // 100 == 1 else BIRLER[g // 100]), "Yüz"] if g // 100 else []) + [ONLAR[(g // 10) % 10], BIRLER[g % 10]]
            kelime = " ".join(x for x in uc if x)
            parca.insert(0, GRUPLAR[i] if (g == 1 and i == 1) else f"{kelime} {GRUPLAR[i]}".strip())
        n //= 1000
        i += 1
    return " ".join(parca)


def yazi_dogru(yazi: str, tutar: float) -> bool:
    lira, kurus = int(round(tutar * 100)) // 100, int(round(tutar * 100)) % 100
    ns = katla(yazi).replace(" ", "")
    beklenen = katla(yaziyla(lira)).replace(" ", "")
    if not ns.startswith(beklenen) or not ns[len(beklenen):].startswith(PARA_KELIME):
        return False
    if kurus:
        return katla(yaziyla(kurus)).replace(" ", "") + "kurus" in ns
    return "kurus" not in ns or "sifirkurus" in ns


def form_oku(yol: Path) -> dict:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: talep formu bulunamadı.")
    f = {a: "" for a in FORM}
    for s in belge.metin_oku(yol).splitlines():
        a, _, d = s.partition(":")
        k = next((k for k, v in FORM.items() if katla(v) == katla(a)), None)
        if k:
            f[k] = d.strip()
    eksik = [FORM[k] for k in ("lehtar", "borclu", "tutar", "para") if not f[k]]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli satır yok: {', '.join(eksik)}")
    try:
        f["tutar_v"] = float(sayi_norm(f["tutar"]))
    except ValueError:
        raise llm.LLMHatasi(f"{yol.name}: tutar okunamadı: '{f['tutar']}'")
    f["vade_v"], f["duzenleme_v"] = tarih_coz(f["vade"]), tarih_coz(f["duzenleme"])
    f["suresiz"] = katla(f["vade"]) == "suresiz"
    f["para_adi"] = PARA_ADI.get(katla(f["para"]), f["para"])
    return f


def para_yaz(v: float) -> str:
    return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def standart_doldur(metin: str, f: dict) -> str:
    deger = {"TUTAR": para_yaz(f["tutar_v"]), "PARA": f["para"], "TUTAR_YAZI": f"{yaziyla(int(f['tutar_v']))} {f['para_adi']}",
             "VADE": f["vade"], "SOZLESME": f["sozlesme"], "TUR": f["tur"]}
    for k, v in deger.items():
        metin = metin.replace("{" + k + "}", v)
    return metin


def gizle(metin: str, f: dict) -> str:
    for ad, kod in sorted(((f["lehtar"], "{LEHTAR}"), (f["borclu"], "{BORCLU}")), key=lambda x: -len(x[0])):
        metin = re.sub(re.escape(ad), kod, metin, flags=re.I)
    return metin


def cumleler(metin: str, on: str) -> dict:
    out = []
    for p in re.split(r"\n\s*\n", metin.strip()):
        p = " ".join(p.split())
        out += [c.strip() for c in re.split(r"(?<=[.!?])\s+(?=[A-ZÇĞİÖŞÜ{])", p) if c.strip()]
    return {f"{on}{i}": c for i, c in enumerate(out, 1)}


def benzerlik(a: str, b: str) -> float:
    # Kesme işaretinden sonraki ek farkı ({LEHTAR}'na / {LEHTAR}'ne) benzerliği düşürmesin
    ek_at = lambda s: re.sub(r"['’]\w+", "", s)  # noqa: E731
    return difflib.SequenceMatcher(None, katla(ek_at(a)), katla(ek_at(b))).ratio()


def alan_kontrol(f: dict, ham: str, standart: dict, talep: dict) -> list[str]:
    s = []
    for alan in ("lehtar", "borclu", "sozlesme"):
        if f[alan] and not metinde_var(f[alan], ham):
            s.append(f"{FORM[alan]}: formdaki '{f[alan]}' talep metninde geçmiyor")
    tutarlar = re.findall(r"(\d{1,3}(?:\.\d{3})*,\d{2})\s*(TL|TRY|USD|EUR|EURO|GBP)?", ham)
    if not tutarlar:
        s.append("Tutar: talep metninde rakamla tutar bulunamadı")
    for t, pb in tutarlar:
        if abs(float(sayi_norm(t)) - f["tutar_v"]) > 0.005 and "%" not in t:
            s.append(f"Tutar: talep metninde formdan farklı tutar var ({t} {pb}); formda {para_yaz(f['tutar_v'])} {f['para']}")
        if pb and PARA_ADI.get(katla(pb), pb) != f["para_adi"]:
            s.append(f"Para birimi: talep metninde {pb}, formda {f['para']}")
    yazilar = re.findall(r"\(([^()]*(?:lira|dolar|avro|euro|sterlin)[^()]*)\)", ham, flags=re.I)
    if not yazilar:
        s.append("Tutar: yazıyla tutar bulunamadı")
    for y in yazilar:
        for v in dict.fromkeys([f["tutar_v"]] + [float(sayi_norm(x)) for x, _ in tutarlar]):
            if not yazi_dogru(y, v):
                s.append(f"Tutar: yazıyla tutar rakamla ({para_yaz(v)}) uyuşmuyor ('{y}'; beklenen '{yaziyla(int(v))} …')")
    gecerli = [c for c in talep.values() if re.search(r"gecerli|suresi", katla(c))]
    tarihler = {tarih_coz(t) for c in gecerli for t in re.findall(r"\d{2}\.\d{2}\.\d{4}", c)} - {None}
    if f["suresiz"]:
        if tarihler:
            s.append(f"Vade: form süresiz, talep metninde geçerlilik tarihi var ({', '.join(f'{d:%d.%m.%Y}' for d in sorted(tarihler))})")
    elif f["vade_v"]:
        if f["vade_v"] not in tarihler:
            bulunan = ", ".join(f"{d:%d.%m.%Y}" for d in sorted(tarihler)) or "yok"
            s.append(f"Vade: talep metnindeki geçerlilik tarihi ({bulunan}) formdaki vadeyle ({f['vade']}) aynı değil")
        if f["duzenleme_v"] and f["vade_v"] <= f["duzenleme_v"]:
            s.append(f"Vade: {f['vade']}, düzenleme tarihinden ({f['duzenleme']}) önce")
    return s


def eksik_standart(standart: dict, talep: dict) -> list[str]:
    return [k for k, v in standart.items() if max((benzerlik(v, t) for t in talep.values()), default=0) < 0.6]


def riskli_mi(metin: str) -> list[str]:
    return [ad for ad, desen in RISKLI if re.search(desen, katla(metin))]


def denetle(t: dict, standart: dict, talep: dict) -> list[str]:
    sorun = []
    gorulen = {}
    std_risk = set(riskli_mi(" ".join(standart.values())))
    izinli = sayilar(" ".join(standart.values()) + " " + " ".join(talep.values()))
    for e in t["eslesmeler"]:
        tid = e["talep"]
        if tid not in talep:
            sorun.append(f"{tid}: talep metninde böyle bir cümle yok")
            continue
        gorulen[tid] = gorulen.get(tid, 0) + 1
        if e["standart"] and e["standart"] not in standart:
            sorun.append(f"{tid}: eşleştirilen standart cümle geçersiz ({e['standart']})")
        if e["durum"] != "Standartta yok" and not e["standart"]:
            sorun.append(f"{tid}: '{e['durum']}' denmiş ama standart cümle numarası yok")
        if e["durum"] == "Aynı" and e["standart"] in standart and benzerlik(talep[tid], standart[e["standart"]]) < 0.9:
            sorun.append(f"{tid}: 'Aynı' denmiş ama {e['standart']} ile metin farklı (benzerlik %{benzerlik(talep[tid], standart[e['standart']]) * 100:.0f})")
        elif e["durum"] == "Aynı" and e["standart"] in standart and sayilar(talep[tid]) != sayilar(standart[e["standart"]]):
            fark = sorted(sayilar(talep[tid]) ^ sayilar(standart[e["standart"]]))
            sorun.append(f"{tid}: 'Aynı' denmiş ama {e['standart']} ile sayı / tarih farklı ({', '.join(fark)})")
        yeni = [r for r in riskli_mi(talep[tid]) if r not in std_risk]
        if yeni and (e["durum"] not in ("Esaslı sapma", "Standartta yok") or e["risk"] == "Düşük"):
            sorun.append(f"{tid}: standartta olmayan riskli hüküm ({', '.join(yeni)}) esaslı sapma ve en az orta risk olarak işaretlenmeli")
        fazla = sorted(sayilar(re.sub(r"\b[ST]\d+\b", " ", e["aciklama"])) - izinli)  # S5, T6 gibi kimlikler sayı değil
        if fazla:
            sorun.append(f"{tid}: açıklamada metinlerde olmayan sayı: {', '.join(fazla)}")
    for tid in talep:
        if tid not in gorulen:
            sorun.append(f"{tid}: talep cümlesi sınıflandırılmamış ('{talep[tid][:50]}')")
        elif gorulen[tid] > 1:
            sorun.append(f"{tid}: birden fazla kez sınıflandırılmış")
    return sorun


def uret(standart: dict, talep: dict, f: dict, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    s = "\n".join(f"{k} | {v}" for k, v in standart.items())
    m = "\n".join(f"{k} | {v}" for k, v in talep.items())
    mesaj = (f"<mektup_turu>{f['tur'] or '-'}</mektup_turu>\n<standart_metin>\n{s}\n</standart_metin>\n<talep_metni>\n{m}\n</talep_metni>"
             + duzeltme_blogu(duzeltme, onceki))
    return llm.json_iste(sistem, mesaj, SEMA)


def ac(metin: str, f: dict) -> str:
    return metin.replace("{LEHTAR}", f["lehtar"]).replace("{BORCLU}", f["borclu"])


def rapor_yaz(cikti: Path, t, standart, talep, f, alanlar, eksik, kalan, tur) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Karşılaştırma"
    baslik(ws, ["No", "Talep Metni", "Standart", "Standart Metin", "Durum", "Risk", "Açıklama"], (6, 60, 9, 60, 14, 9, 50))
    renk = {"Aynı": YESIL, "İfade farkı": SARI, "Esaslı sapma": KIRMIZI, "Standartta yok": KIRMIZI}
    for e in sorted(t["eslesmeler"], key=lambda x: int(re.sub(r"\D", "", x["talep"]) or 0)):
        ws.append([e["talep"], ac(talep.get(e["talep"], "?"), f), e["standart"], ac(standart.get(e["standart"], ""), f), e["durum"], e["risk"], e["aciklama"]])
        boya(ws, (5,), renk[e["durum"]])
        boya(ws, (7,), AI_DOLGU)
    for k in eksik:
        ws.append(["", "(talep metninde yok)", k, ac(standart[k], f), "Eksik standart hüküm", "", "Kod: standart cümle talep metninde bulunamadı"])
        boya(ws, (5,), KIRMIZI)
    son_not(ws, "Genel değerlendirme (model): " + t["genel_degerlendirme"] + "\nÖnemli: Ön kontroldür. Sapmaların kabulü bankanın yetkili birimlerinin "
                "(hukuk, kredi) onayına tabidir.", kalan, tur)
    hizala(ws)
    w = wb.create_sheet("Alan Kontrol")
    baslik(w, ["Kontrol", "Sonuç"], (30, 100))
    for ad in ("Lehtar", "Borçlu", "Sözleşme", "Tutar", "Para birimi", "Vade"):
        ilgili = [x.split(": ", 1)[1] for x in alanlar if x.startswith(ad + ":")]
        w.append([ad, "; ".join(ilgili) or "Formla uyumlu"])
        boya(w, (2,), KIRMIZI if ilgili else YESIL)
    hizala(w)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    sap = [e for e in t["eslesmeler"] if e["durum"] in ("Esaslı sapma", "Standartta yok")]
    metin = "\n".join([f"TEMİNAT MEKTUBU METİN KONTROLÜ · {f['lehtar']} · {para_yaz(f['tutar_v'])} {f['para']}", "",
                       f"Alan kontrolü: {len(alanlar)} sorun"] + [f"  - {x}" for x in alanlar] +
                      ["", f"Esaslı sapma / standartta olmayan hüküm: {len(sap)}"] +
                      [f"  - {e['talep']} [{e['risk']}] {ac(talep.get(e['talep'], ''), f)[:120]} → {e['aciklama']}" for e in sap] +
                      ["", f"Talep metninde olmayan standart hüküm: {', '.join(eksik) or 'yok'}", "", t["genel_degerlendirme"]])
    cikti.with_suffix(".txt").write_text(metin, encoding="utf-8")
    return metin


def calistir(form: Path, standart_yol: Path, metin_yol: Path, cikti: Path, tur: int = 2, evet: bool = False) -> dict:
    f = form_oku(form)
    for yol in (standart_yol, metin_yol):
        if not yol.exists():
            raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    ham = belge.metin_oku(metin_yol)
    standart = cumleler(gizle(standart_doldur(belge.metin_oku(standart_yol), f), f), "S")
    talep = cumleler(gizle(ham, f), "T")
    alanlar = alan_kontrol(f, ham, standart, talep)
    eksik = eksik_standart(standart, talep)
    print(f"[OK] {len(standart)} standart · {len(talep)} talep cümlesi · {len(alanlar)} alan sorunu · eksik standart: {', '.join(eksik) or 'yok'}")
    llm.onay_al(f"Standart metin ve talep metni ({len(talep)} cümle) gönderilecek; lehtar ve borçlu unvanları kodlanır.", evet)
    t, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(standart, talep, f, d, o), lambda x: denetle(x, standart, talep), tur)
    metin = rapor_yaz(cikti, t, standart, talep, f, alanlar, eksik, kalan, yapilan)
    return {"form": f, "alanlar": alanlar, "eksik": eksik, "taslak": t, "kalan": kalan, "tur": yapilan, "metin": metin, "talep": talep}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    p = argparse.ArgumentParser(description="Teminat mektubu talep metnini banka standart metni ve talep formuyla karşılaştırıp sapmaları işaretler.")
    v = BURASI / "ornek_veri"
    p.add_argument("--form", type=Path, default=v / "talep_formu.txt", help="Talep formu (.txt/.docx): Lehtar, Borçlu, Tür, Tutar, Para Birimi, Vade, Düzenleme Tarihi, Sözleşme")
    p.add_argument("--standart", type=Path, default=v / "standart_metin.txt", help="Banka standart metni ({LEHTAR}, {BORCLU}, {TUTAR}, {PARA}, {TUTAR_YAZI}, {VADE}, {SOZLESME})")
    p.add_argument("--metin", type=Path, default=v / "talep_metni.txt", help="Lehtarın / müşterinin istediği mektup metni (.txt/.docx/.pdf)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "teminat_kontrol.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.form, a.standart, a.metin, a.cikti, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    for x in s["alanlar"]:
        print(f"[!] {x}")
    sap = sum(e["durum"] in ("Esaslı sapma", "Standartta yok") for e in s["taslak"]["eslesmeler"])
    print(f"[OK] {sap} esaslı sapma / standartta olmayan hüküm")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
