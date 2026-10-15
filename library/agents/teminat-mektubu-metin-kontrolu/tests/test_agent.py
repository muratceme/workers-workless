"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import belge  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

V = KLASOR / "ornek_veri"
FORM, STD, MET = V / "talep_formu.txt", V / "standart_metin.txt", V / "talep_metni.txt"
MESAJLAR = []


def e(t, s, durum, risk, aciklama):
    return {"talep": t, "standart": s, "durum": durum, "risk": risk, "aciklama": aciklama}


def cevap(dogru=True) -> dict:
    t = {"eslesmeler": [
        e("T1", "S1", "İfade farkı", "Düşük", "Yalnız ek farkı."),
        e("T2", "S2", "Aynı", "Düşük", "Standartla aynı."),
        e("T3", "S3", "Aynı", "Düşük", "Standartla aynı."),
        e("T4", "", "Standartta yok", "Yüksek", "Standartta faiz hükmü yok; banka 3 iş günü sonrası için yıllık %48 temerrüt faizi ödemeyi taahhüt eder."),
        e("T5", "S4", "Esaslı sapma", "Yüksek", "Standartta geçerlilik 31.12.2027, talepte 31.12.2028."),
        e("T6", "", "Standartta yok", "Yüksek", "Kendiliğinden uzama hükmü bankanın yükümlülüğünü süresiz hâle getirir."),
        e("T7", "", "Standartta yok", "Yüksek", "Devir ve temlik standartta yok."),
        e("T8", "S6", "Aynı", "Düşük", "Standartla aynı.")],
        "genel_degerlendirme": "T4, T6 ve T7 çıkarılmadan ve geçerlilik tarihi forma uygun hâle getirilmeden metin kabul edilmemeli."}
    if not dogru:
        t["eslesmeler"][4] = e("T5", "S4", "Aynı", "Düşük", "Standartla aynı.")
        t["eslesmeler"][5] = e("T6", "S5", "İfade farkı", "Düşük", "Süreyle ilgili ifade farkı; 2 yıl uzar.")
        del t["eslesmeler"][6]
    return t


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return cevap("<duzeltme>" in kullanici)


def metinler():
    f = agent.form_oku(FORM)
    std = agent.cumleler(agent.gizle(agent.standart_doldur(belge.metin_oku(STD), f), f), "S")
    ham = belge.metin_oku(MET)
    return f, std, ham, agent.cumleler(agent.gizle(ham, f), "T")


class KodTesti(unittest.TestCase):
    def test_yaziyla(self):
        self.assertEqual([agent.yaziyla(n) for n in (1000, 1100, 2450000, 1000000, 101)],
                         ["Bin", "Bin Yüz", "İki Milyon Dört Yüz Elli Bin", "Bir Milyon", "Yüz Bir"])
        self.assertTrue(agent.yazi_dogru("İkimilyondörtyüzellibin Türk Lirası", 2450000))
        self.assertTrue(agent.yazi_dogru("On Beş Bin İki Yüz Elli Türk Lirası Elli Kuruş", 15250.5))
        self.assertFalse(agent.yazi_dogru("İki Milyon Dört Yüz Bin Türk Lirası", 2450000))
        self.assertFalse(agent.yazi_dogru("İki Milyon Dört Yüz Elli Bin Beş Türk Lirası", 2450000))

    def test_alan_kontrol(self):
        f, std, ham, tal = metinler()
        self.assertEqual((len(std), len(tal)), (6, 8))
        self.assertEqual(agent.alan_kontrol(f, ham, std, tal), ["Vade: talep metnindeki geçerlilik tarihi (31.12.2028) formdaki vadeyle (31.12.2027) aynı değil"])
        self.assertEqual(agent.eksik_standart(std, tal), ["S5"])
        bozuk = ham.replace("2.450.000,00 TL (İki Milyon Dört Yüz Elli Bin", "2.540.000,00 USD (İki Milyon Dört Yüz Elli Bin").replace("Örnek Yapı Taahhüt A.Ş.", "Örnek Yapı A.Ş.")
        s = agent.alan_kontrol(f, bozuk, std, agent.cumleler(agent.gizle(bozuk, f), "T"))
        self.assertTrue(any(x.startswith("Lehtar") for x in s) and any("2.540.000,00" in x for x in s) and any("Para birimi" in x for x in s)
                        and any("yazıyla" in x for x in s), s)

    def test_denetim(self):
        f, std, _, tal = metinler()
        self.assertEqual(agent.denetle(cevap(), std, tal), [])
        s = agent.denetle(cevap(False), std, tal)
        for parca in ("T5: 'Aynı' denmiş ama S4 ile sayı", "T6: standartta olmayan riskli", "T6: açıklamada", "T7: talep cümlesi sınıflandırılmamış"):
            self.assertTrue(any(parca in x for x in s), (parca, s))
        self.assertEqual(len(s), 4, s)


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "k.xlsx"
            s = agent.calistir(FORM, STD, MET, cikti, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            for gizli in ("Örnek Yapı", "Deneme Mühendislik"):
                self.assertNotIn(gizli, "".join(MESAJLAR))
            ws = load_workbook(cikti)["Karşılaştırma"]
            self.assertEqual((ws["A5"].value, ws["E5"].value), ("T4", "Standartta yok"))
            self.assertIn("Deneme Mühendislik Ltd. Şti.", ws["B3"].value)
            self.assertEqual((ws["C10"].value, ws["E10"].value), ("S5", "Eksik standart hüküm"))
            self.assertEqual(load_workbook(cikti)["Alan Kontrol"]["B7"].value.split(":")[0], "talep metnindeki geçerlilik tarihi (31.12.2028) formdaki vadeyle (31.12.2027) aynı değil".split(":")[0])
            self.assertIn("Esaslı sapma / standartta olmayan hüküm: 4", s["metin"])

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(FORM, STD, MET, Path(tmp) / "k.xlsx", tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 4)
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--evet", "--metin", str(Path(tmp) / "yok.txt")]), 1)


if __name__ == "__main__":
    unittest.main()
