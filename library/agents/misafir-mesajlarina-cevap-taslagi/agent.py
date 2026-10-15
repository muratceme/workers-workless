"""
Misafir Mesajlarına Cevap Taslağı — Workers / Workless AI Agent
Ön Büro › Resepsiyonist

1. Kod otel bilgi dosyasını numaralı bilgilere (B1, B2 …) ayırır. Misafir adı ve rezervasyon numarası modele
   gönderilmez; cevapta ad yerine {AD} yer tutucusu kullanılır, kod doldurur.
2. Model her mesaj için misafirin dilinde cevap taslağı yazar, kullandığı otel bilgilerini (B no) belirtir; bilgi
   dosyasında cevabı olmayan soruları ve insan onayı gereken durumları (iptal, iade, şikâyet, ücret istisnası, özel
   talep) işaretler.
3. Kod denetler:
   - Belirtilen B numaraları var mı
   - Cevaptaki her sayı, saat ve ücret kullanılan bilgilerde ya da misafirin mesajında geçiyor mu
   - "Ücretsiz / garanti" gibi ifadeler dayanaksız mı, "onaylandı / confirmed" gibi kesin onay verilmiş mi
   - İptal / iade / şikâyet içeren mesaj ya da cevaplanamayan soru insan onayına işaretlenmiş mi
   - {AD} yer tutucusu var mı
   Sorun varsa yalnız o mesaj modele geri gönderilir (en fazla --tur kez).
4. Taslaklar gönderilmez; Excel'de "Onay" sütunuyla listelenir.

Kullanım:
    python agent.py                                              # örnek mesajlar + otel bilgileri
    python agent.py --girdi mesajlar.xlsx --bilgi otel_bilgileri.txt
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
PAKET = 15
SEMA = {
    "type": "object",
    "properties": {"cevaplar": {"type": "array", "items": {"type": "object", "properties": {
        "id": S, "dil": S, "konu_ozeti": S, "cevap": S, "kullanilan_bilgiler": {"type": "array", "items": S},
        "cevaplanamayan": {"type": "array", "items": S}, "insan_gerekli": {"type": "boolean"}, "insan_notu": S},
        "required": ["id", "dil", "konu_ozeti", "cevap", "kullanilan_bilgiler", "cevaplanamayan", "insan_gerekli", "insan_notu"], "additionalProperties": False}}},
    "required": ["cevaplar"],
    "additionalProperties": False,
}
TAAHHUT = re.compile(r"(ücretsiz|ucretsiz|garanti|free of charge|\bfree\b|guarantee|kostenlos|garantiert|бесплатн|гарант)", re.I)
TAAHHUT_BILGI = re.compile(r"(ucretsiz|garanti)")
ONAY = re.compile(r"(onaylandı|onayladık|onaylanmıştır|\bconfirmed\b|bestätigt|подтвержден)", re.I)
INSAN = re.compile(r"(iade|iptal|şikayet|şikâyet|refund|cancel|complain|erstatt|stornier|beschwerde|возврат|отмен|жалоб)", re.I)


def bilgileri_oku(yol: Path) -> dict[str, str]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: otel bilgi dosyası bulunamadı.")
    out, sira = {}, 0
    for s in belge.metin_oku(yol).splitlines():
        s = s.strip()
        if not s:
            continue
        m = re.match(r"^(B\d+)\s*[:.)-]\s*(.+)$", s)
        if m:
            out[m.group(1)] = m.group(2)
        else:
            sira += 1
            while f"B{sira}" in out:
                sira += 1
            out[f"B{sira}"] = s
    return out


def mesajlari_oku(yol: Path) -> list[dict]:
    r = kayitlar(yol, {"id": ("Mesaj No", "No"), "zaman": ("Tarih", "Zaman"), "kanal": ("Kanal",), "ad": ("Misafir Adı", "Misafir", "Ad Soyad"),
                       "rez": ("Rezervasyon No", "Rezervasyon"), "mesaj": ("Mesaj", "Metin", "E-posta")}, ("mesaj",))
    return [{"id": str(x["id"] or f"MS{i}"), "zaman": zaman_coz(x["zaman"]), "kanal": str(x["kanal"]), "ad": str(x["ad"]), "rez": str(x["rez"]),
             "mesaj": str(x["mesaj"]).strip()} for i, x in enumerate(r, 1) if str(x["mesaj"]).strip()]


def gonderilecek(m: dict) -> str:
    t = m["mesaj"]
    for x in (m["rez"], m["ad"], *m["ad"].split()):
        if x and len(x) >= 3:
            t = re.sub(re.escape(x), "[REZ]" if x == m["rez"] else "[MİSAFİR]", t, flags=re.I)
    return llm.maskele(t)


def denetle(mesaj: list[dict], sonuc: dict, bilgi: dict[str, str]) -> list[str]:
    sorun = []
    for m in mesaj:
        c = sonuc.get(m["id"])
        if not c:
            sorun.append(f"{m['id']}: cevap yazılmadı")
            continue
        yok = [b for b in c["kullanilan_bilgiler"] if b not in bilgi]
        if yok:
            sorun.append(f"{m['id']}: bilgi dosyasında olmayan bilgi numarası {yok}")
        dayanak = " ".join(bilgi.get(b, "") for b in c["kullanilan_bilgiler"])
        fazla = sorted(sayilar(c["cevap"]) - sayilar(dayanak + " " + gonderilecek(m)))
        if fazla:
            sorun.append(f"{m['id']}: kullanılan bilgilerde ve mesajda olmayan sayı / saat / ücret: {', '.join(fazla)} (bilgi numarasını ekleyin ya da çıkarın)")
        if TAAHHUT.search(c["cevap"]) and not TAAHHUT_BILGI.search(katla(dayanak)):
            sorun.append(f"{m['id']}: 'ücretsiz / garanti' ifadesi kullanılan bilgilere dayanmıyor")
        if ONAY.search(c["cevap"]):
            sorun.append(f"{m['id']}: cevapta kesin onay ifadesi var; taslak talebi onaylayamaz, 'talebinizi ilettik' gibi yazın")
        if (INSAN.search(m["mesaj"]) or c["cevaplanamayan"]) and not c["insan_gerekli"]:
            sorun.append(f"{m['id']}: iptal / iade / şikâyet içeriyor ya da cevaplanamayan soru var; insan_gerekli true olmalı")
        if "{AD}" not in c["cevap"]:
            sorun.append(f"{m['id']}: hitapta {{AD}} yer tutucusu yok")
    return sorun


def uret(mesaj: list[dict], bilgi: dict[str, str], duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    sonuc, hedef = {}, mesaj
    if duzeltme is not None:  # düzeltmede yalnız sorunlu mesajlar yeniden yazılır
        sorunlu = {s.split(":")[0] for s in duzeltme}
        sonuc.update({a: b for a, b in (onceki or {}).items() if a not in sorunlu})
        hedef = [m for m in mesaj if m["id"] in sorunlu]
    b = "\n".join(f"{k}: {v}" for k, v in bilgi.items())
    for j in range(0, len(hedef), PAKET):
        parca = hedef[j:j + PAKET]
        ids = {m["id"] for m in parca}
        satir = "\n".join(f'<mesaj id="{m["id"]}" kanal="{m["kanal"]}">{gonderilecek(m)}</mesaj>' for m in parca)
        d = [s for s in (duzeltme or []) if s.split(":")[0] in ids]
        mesaj_metni = f"<otel_bilgileri>\n{b}\n</otel_bilgileri>\n<mesajlar>\n{satir}\n</mesajlar>" + duzeltme_blogu(d, [v for a, v in (onceki or {}).items() if a in ids])
        for c in llm.json_iste(sistem, mesaj_metni, SEMA).get("cevaplar", []):
            if c.get("id") in ids:
                sonuc[c["id"]] = c
    return sonuc


def doldur(cevap: str, ad: str, dil: str) -> str:
    if ad:
        return cevap.replace("{AD}", ad)
    varsayilan = {"tr": "Değerli misafirimiz", "en": "Guest", "de": "Gast", "ru": "гость"}
    return cevap.replace("{AD}", varsayilan.get(dil.lower()[:2], "Guest"))


def rapor_yaz(cikti: Path, mesaj, sonuc, bilgi, kalan, tur) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Cevap Taslakları"
    baslik(ws, ["Mesaj No", "Zaman", "Kanal", "Misafir", "Rezervasyon", "Mesaj", "Dil", "Konu (TR)", "Cevap Taslağı", "Kullanılan Bilgiler", "Cevaplanamayan",
                "İnsan Onayı", "Not", "Onay"], (8, 15, 12, 16, 11, 50, 5, 30, 70, 14, 30, 10, 36, 8))
    for m in sorted(mesaj, key=lambda m: not sonuc.get(m["id"], {}).get("insan_gerekli", True)):
        c = sonuc.get(m["id"])
        if not c:
            ws.append([m["id"], m["zaman"], m["kanal"], m["ad"], m["rez"], m["mesaj"], "", "", "AI yanıt vermedi"])
            continue
        ws.append([m["id"], m["zaman"], m["kanal"], m["ad"], m["rez"], m["mesaj"], c["dil"], c["konu_ozeti"], doldur(c["cevap"], m["ad"], c["dil"]),
                   ", ".join(c["kullanilan_bilgiler"]), "; ".join(c["cevaplanamayan"]), "GEREKLİ" if c["insan_gerekli"] else "", c["insan_notu"], ""])
        ws.cell(ws.max_row, 2).number_format = "DD.MM.YYYY HH:MM"
        boya(ws, [8, 9], AI_DOLGU)
        if c["insan_gerekli"]:
            boya(ws, [12], KIRMIZI)
        boya(ws, [14], SARI)
    ws.freeze_panes = "B2"
    son_not(ws, "Önemli: Taslaklar gönderilmedi. İnsan onayı gerekenleri ilgili yetkiliye iletin; diğerlerini okuyup onayladıktan sonra gönderin.", kalan, tur)
    hizala(ws)
    b = wb.create_sheet("Otel Bilgileri")
    baslik(b, ["No", "Bilgi", "Kullanıldığı Mesajlar"], (6, 100, 24))
    for k, v in bilgi.items():
        b.append([k, v, ", ".join(m for m, c in sonuc.items() if k in c["kullanilan_bilgiler"])])
    hizala(b)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(girdi: Path, bilgi_yolu: Path, cikti: Path, tur: int = 2, evet: bool = False) -> dict:
    mesaj = mesajlari_oku(girdi)
    bilgi = bilgileri_oku(bilgi_yolu)
    if not mesaj:
        raise llm.LLMHatasi(f"{girdi.name}: mesaj yok.")
    print(f"[OK] {len(mesaj)} mesaj · {len(bilgi)} otel bilgisi")
    llm.onay_al(f"{len(mesaj)} misafir mesajı (ad ve rezervasyon no çıkarılmış; telefon / e-posta maskeli) ve otel bilgileri gönderilecek.", evet)
    sonuc, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(mesaj, bilgi, d, o), lambda x: denetle(mesaj, x, bilgi), tur)
    rapor_yaz(cikti, mesaj, sonuc, bilgi, kalan, yapilan)
    return {"mesajlar": mesaj, "bilgi": bilgi, "sonuc": sonuc, "kalan": kalan, "tur": yapilan}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    p = argparse.ArgumentParser(description="Misafir mesajlarına otel bilgilerine dayanan çok dilli cevap taslakları hazırlar (göndermez).")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "mesajlar.csv",
                   help="Mesajlar (.xlsx/.csv): Mesaj; Mesaj No, Tarih, Kanal, Misafir Adı, Rezervasyon No isteğe bağlı")
    p.add_argument("--bilgi", type=Path, default=BURASI / "ornek_veri" / "otel_bilgileri.txt",
                   help="Otel bilgileri (.txt/.md/.docx): her satırda bir bilgi, isteğe bağlı 'B1:' numarasıyla")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "misafir_cevaplari.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.bilgi, a.cikti, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    insan = [k for k, c in s["sonuc"].items() if c["insan_gerekli"]]
    print(f"[OK] {len(s['sonuc'])}/{len(s['mesajlar'])} taslak · insan onayı gereken: {', '.join(insan) or 'yok'}")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
