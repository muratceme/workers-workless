"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

MESAJ = KLASOR / "ornek_veri" / "mesajlar.csv"
BILGI = KLASOR / "ornek_veri" / "otel_bilgileri.txt"
MESAJLAR = []
DOGRU = {
    "MS-01": ("en", "Dear {AD}, thank you for your message. Our check-in time is 14:00; early check-in depends on room availability and cannot be guaranteed "
                    "in advance, but you can leave your luggage at reception free of charge. Airport transfer is 45 EUR one way (up to 3 people) and must be "
                    "requested at least 24 hours in advance. Kind regards, Reception", ["B1", "B5"], False),
    "MS-02": ("tr", "Sayın {AD}, geç çıkış müsaitliğe göre en geç 16:00'ya kadar verilebilir; ücreti 50 EUR'dur. Talebinizi giriş sırasında resepsiyona "
                    "iletmenizi rica ederiz. Standart çıkış saatimiz 12:00'dir.", ["B2", "B3"], False),
    "MS-03": ("de", "Sehr geehrte/r {AD}, leider können wir keine Haustiere aufnehmen (ausgenommen Blindenführhunde).", ["B7"], False),
    "MS-04": ("tr", "Sayın {AD}, geçmiş olsun. İade edilemez tarifede iptal yapılamamaktadır; talebinizi rezervasyon müdürümüzün değerlendirmesine "
                    "ilettik, size en kısa sürede dönüş yapılacaktır.", ["B9"], True),
    "MS-05": ("ru", "Здравствуйте, {AD}! Завтрак подается с 07:00 до 10:30. Открытый бассейн работает с мая по октябрь с 08:00 до 19:00.", ["B4", "B8"], False),
}


def cevaplar(dogru=True) -> dict:
    d = {i: {"id": i, "dil": dl, "konu_ozeti": "özet", "cevap": c, "kullanilan_bilgiler": b, "cevaplanamayan": [], "insan_gerekli": ins, "insan_notu": ""}
         for i, (dl, c, b, ins) in DOGRU.items()}
    if not dogru:
        d["MS-01"]["cevap"] = d["MS-01"]["cevap"].replace("45 EUR", "40 EUR")
        d["MS-04"]["insan_gerekli"] = False
        d["MS-02"]["cevap"] += " Talebiniz onaylandı."
    return d


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    d = cevaplar("<duzeltme>" in kullanici)
    return {"cevaplar": [d[i] for i in re.findall(r'<mesaj id="([^"]+)"', kullanici)]}


class KodTesti(unittest.TestCase):
    def test_okuma(self):
        b = agent.bilgileri_oku(BILGI)
        self.assertEqual((len(b), b["B4"][:14]), (10, "Açık büfe kahv"))
        m = agent.mesajlari_oku(MESAJ)
        self.assertEqual(len(m), 5)
        self.assertEqual(agent.doldur("Dear {AD},", "", "en"), "Dear Guest,")
        self.assertEqual(agent.doldur("Sayın {AD},", "Ayşe Demir", "tr"), "Sayın Ayşe Demir,")

    def test_denetim(self):
        m, b = agent.mesajlari_oku(MESAJ), agent.bilgileri_oku(BILGI)
        self.assertEqual(agent.denetle(m, cevaplar(), b), [])
        s = agent.denetle(m, cevaplar(False), b)
        self.assertEqual(len(s), 3)
        self.assertTrue(any("MS-01" in x and "40" in x for x in s) and any("MS-04" in x and "insan_gerekli" in x for x in s) and any("MS-02" in x and "onay" in x for x in s))
        d = cevaplar()
        d["MS-03"]["kullanilan_bilgiler"] = ["B7", "B12"]
        d["MS-03"]["cevap"] = d["MS-03"]["cevap"].replace("{AD}", "Frau Müller") + " Parkplatz kostenlos."
        self.assertEqual(len(agent.denetle(m, d, b)), 3)


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "m.xlsx"
            s = agent.calistir(MESAJ, BILGI, cikti, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            for gizli in ("John", "Smith", "RZ-77", "Kemal", "Müller"):
                self.assertNotIn(gizli, MESAJLAR[0])
            self.assertEqual(re.findall(r'<mesaj id="([^"]+)"', MESAJLAR[1]), ["MS-01", "MS-02", "MS-04"])
            ws = load_workbook(cikti)["Cevap Taslakları"]
            self.assertEqual((ws.cell(2, 1).value, ws.cell(2, 12).value), ("MS-04", "GEREKLİ"))
            self.assertTrue(ws.cell(3, 9).value.startswith("Dear John Smith,"))

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(MESAJ, BILGI, Path(tmp) / "m.xlsx", tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 3)
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--bilgi", str(BILGI.with_name("yok.txt")), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
