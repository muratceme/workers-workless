"""
Risk Değerlendirme Özeti — Workers / Workless AI Agent
Sigorta › Teknik (Risk Kabul / Underwriting) › Teknik Uzman (Underwriter)

1. Klasördeki risk inceleme (survey) raporu ve ekleri okunur (.txt/.pdf/.docx); isteğe bağlı sigorta bedelleri
   tablosu (Kalem, Bedel).
2. Kod kontrol eder: kontrol listesindeki konular (yapı tarzı, çatı/cephe, yangın algılama ve söndürme, elektrik
   tesisatı, yanıcı madde, sıcak çalışma, itfaiye, yıldırımdan korunma, deprem/zemin, raf ankrajı, sel, hırsızlık
   önlemleri, İSG, hasar geçmişi) raporda var mı; koruma önlemlerinin "yok / bulunmamaktadır / eksik" diye geçtiği
   cümleler (olumsuz gözlem adayları); emtia bedeli raporda yazan maksimum stok değerinin altında mı; rapordaki
   numaralı öneriler.
3. Sigortalı unvanı, risk adresi, görüşülen kişiler ve --gizle ile verilen adlar takma adlarla; telefon, e-posta,
   TCKN ve IBAN maskelenir. Model yangın, deprem, sel, hırsızlık ve sorumluluk risklerini rapordan birebir
   alıntılarla özetler; iyileştirme önerilerini ve kabul önerisini (taslak) yazar.
4. Kod modelin çıktısını denetler: alıntılar raporda birebir geçiyor mu, rapordaki her öneri özette var mı, "Kabul"
   önerisi yüksek risk veya kabul öncesi iyileştirmeyle çelişiyor mu; girdilerde olmayan sayılar işaretlenir.
5. Çıktı: Markdown risk özeti + Excel (Özet + 'Underwriter Kararı', Riskler, İyileştirmeler + 'Takip',
   Kontrol Listesi, Kontroller).

Kullanım:
    python agent.py                                          # örnek: kurgusal plastik ambalaj fabrikası
    python agent.py --girdi ./risk_klasoru --gizle "Görüşülen Kişi"
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
SIFIR = Decimal(0)
RISK_TURLERI = ["Yangın", "Deprem", "Sel / su baskını", "Hırsızlık", "Sorumluluk", "Diğer"]
# (risk türü, konu, katlanmış anahtar kelimeler)
KONTROL_LISTESI = [
    ("Yangın", "Yapı tarzı / taşıyıcı sistem", ("betonarme", "celik konstruksiyon", "yigma", "ahsap", "yapi tarzi", "prefabrik")),
    ("Yangın", "Çatı ve cephe malzemesi", ("cati", "cephe", "sandvic panel")),
    ("Yangın", "Yangın algılama", ("dedektor", "algilama", "yangin alarm")),
    ("Yangın", "Otomatik söndürme / sprinkler", ("sprinkler", "yagmurlama", "otomatik sondurme")),
    ("Yangın", "Yangın dolabı, hidrant, tüp", ("yangin dolab", "hidrant", "yangin tup", "sondurucu")),
    ("Yangın", "Yangın suyu ve pompa", ("yangin suyu", "yangin pompa", "su deposu")),
    ("Yangın", "Elektrik tesisatı / termal ölçüm", ("elektrik", "termal")),
    ("Yangın", "Yanıcı ve parlayıcı maddeler", ("yanici", "parlayici", "solvent", "tiner", "lpg", "akaryakit")),
    ("Yangın", "Sıcak çalışma", ("sicak calisma", "kaynak")),
    ("Yangın", "Yangın bölümleri / yangın duvarı", ("yangin duvar", "yangin bolum", "yangin kapi")),
    ("Yangın", "İtfaiye mesafesi", ("itfaiye",)),
    ("Yangın", "Yıldırımdan korunma", ("paratoner", "yildirim")),
    ("Yangın", "Düzen ve temizlik", ("duzen", "temizlik")),
    ("Deprem", "Bina yaşı / inşa yılı", ("insa", "yapim yili", "bina yasi", "yilinda")),
    ("Deprem", "Zemin etüdü / deprem tehlikesi", ("zemin", "deprem")),
    ("Deprem", "Raf ve ekipman sabitleme", ("ankraj", "sabitle")),
    ("Deprem", "Ruhsat / iskân", ("iskan", "ruhsat")),
    ("Sel / su baskını", "Dere / su kaynağına mesafe", ("dereye", "dereden", "derenin", "deresi", "nehir", "irmak", "deniz", "golet")),
    ("Sel / su baskını", "Bodrum / drenaj / su baskını geçmişi", ("bodrum", "drenaj", "su baskini", "sel")),
    ("Hırsızlık", "Çevre güvenliği", ("cevre duvar", "tel cit", "cit", "duvar ile")),
    ("Hırsızlık", "Güvenlik görevlisi", ("guvenlik gorevlisi", "bekci", "ozel guvenlik")),
    ("Hırsızlık", "Kamera / alarm", ("kamera", "alarm")),
    ("Sorumluluk", "Çalışan sayısı ve İSG", ("calisan", "is sagligi", "isg", "osgb")),
    ("Sorumluluk", "Üçüncü kişi / ziyaretçi", ("ziyaretci", "ucuncu kisi", "musteri giris")),
    ("Diğer", "Hasar geçmişi", ("hasar gecmis", "hasar kayd", "hasar frekans", "hasarsiz", "odenmistir", "hasar tutari")),
]
KORUMA = ("sprinkler", "dedektor", "alarm", "yangin duvar", "hortum", "izin", "termal", "ankraj", "zemin etud", "paratoner", "kamera",
          "guvenlik", "hidrant", "yangin dolab", "yangin tup", "yangin kapi", "pompa", "topraklama", "kilit")
ONERI_RE = re.compile(r"^\s*(?:R|Ö|Öneri\s*)\d+[.)\-:]\s*")
OLUMSUZ = re.compile(r"\b(bulunmamaktadir|bulunmuyor|yoktur|yok|mevcut degil|degildir|yapilmamistir|yapilmamis|eksiktir|eksik|arizali|"
                     r"calismiyor|calismamaktadir|yetersiz|gorulmemistir|uygun degil)\b")


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def katla(s) -> str:
    s = kucuk(s)
    for a, b in zip("ıöüşğçâîû", "iousgcaiu"):
        s = s.replace(a, b)
    return re.sub(r"[^a-z0-9%]+", " ", s).strip()


def duz(s: str) -> str:
    return " ".join(katla(s).split())


def para(x) -> Decimal | None:
    if x in (None, "", "—", "-"):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("₺", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tl(x) -> str:
    return "—" if x is None else f"{x:,.0f}".replace(",", ".")


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = belge.txt_oku(yol)
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def cumleler(metin: str) -> list[str]:
    return [c.strip() for c in re.split(r"(?<=[.!?])\s+|\n", metin) if len(c.strip()) > 3]


def anahtarli(metin: str, *etiketler: str) -> str | None:
    for et in etiketler:
        m = re.search(r"^\s*" + re.escape(et) + r"\s*:\s*(.+)$", metin or "", re.I | re.M)
        if m:
            return m.group(1).strip()
    return None


# ----------------------------------------------------------------------------
# Girdiler
# ----------------------------------------------------------------------------

@dataclass
class Dosya:
    ad: str
    metin: str


def dosyalari_oku(klasor: Path) -> tuple[list[Dosya], list[dict], list[str]]:
    raporlar, bedeller, atlanan = [], [], []
    for yol in sorted(klasor.iterdir()):
        if yol.suffix.lower() in {".csv", ".xlsx", ".xlsm"}:
            try:
                s = tablo_oku(yol)
            except Exception as h:  # noqa: BLE001
                atlanan.append(f"{yol.name} ({h})")
                continue
            bi = next((i for i, r in enumerate(s[:10]) if any(katla(c) in ("bedel", "sigorta bedeli", "tutar") for c in r)), None)
            if bi is None:
                atlanan.append(f"{yol.name} (Bedel sütunu yok)")
                continue
            b = [katla(c) for c in s[bi]]
            ib = next(j for j, x in enumerate(b) if x in ("bedel", "sigorta bedeli", "tutar"))
            ik = next((j for j, x in enumerate(b) if x in ("kalem", "teminat", "kiymet", "sigortalanan", "aciklama")), 0)
            for r in s[bi + 1:]:
                if ik < len(r) and r[ik] and not katla(r[ik]).startswith("toplam"):
                    bedeller.append({"kalem": str(r[ik]).strip(), "bedel": para(r[ib]) if ib < len(r) else None})
            continue
        if yol.suffix.lower() not in belge.DESTEKLENEN:
            atlanan.append(yol.name)
            continue
        try:
            raporlar.append(Dosya(yol.name, belge.metin_oku(yol)))
        except belge.BelgeHatasi as h:
            atlanan.append(f"{yol.name} ({h})")
    return raporlar, bedeller, atlanan


# ----------------------------------------------------------------------------
# Kontroller
# ----------------------------------------------------------------------------

def kontrol_et(raporlar: list[Dosya], bedeller: list[dict]) -> dict:
    bulgular = []

    def b(onem, aciklama, kaynak=""):
        bulgular.append((onem, aciklama, kaynak))

    tum = "\n".join(d.metin for d in raporlar)
    govde = "\n".join(x for x in tum.splitlines() if not ONERI_RE.match(x) and not (x.strip().upper() == x.strip() and len(x.strip()) < 60))
    cum = cumleler(govde)
    liste = []
    for tur, konu, anahtarlar in KONTROL_LISTESI:
        bulunan = next((c for c in cum if any(f" {a}" in f" {katla(c)}" for a in anahtarlar)), None)
        liste.append({"tur": tur, "konu": konu, "var": bulunan is not None, "cumle": bulunan or ""})
        if bulunan is None:
            b("orta" if tur in ("Yangın", "Deprem") else "bilgi", f"Raporda bilgi yok: {tur} — {konu}", "kontrol listesi")
    adaylar = []
    for c in cum:
        k = katla(c)
        if any(x in k for x in KORUMA) and OLUMSUZ.search(k) and not ONERI_RE.match(c):
            adaylar.append(c)
            b("orta", f"Olumsuz gözlem: {c[:180]}", "rapor")
    oneriler = [x.strip() for x in tum.splitlines() if ONERI_RE.match(x)]
    if not oneriler:
        b("bilgi", "Raporda numaralı öneri bölümü bulunamadı (R1., Öneri 1. gibi)", "rapor")
    bedel = {"toplam": sum((x["bedel"] or SIFIR for x in bedeller), SIFIR)}
    stok = [para(m.group(1)) for m in re.finditer(r"(?:maksimum|azami|en yüksek|en yuksek)\s+stok[^\n.]{0,40}?([\d.]{5,}(?:,\d+)?)\s*TL", tum, re.I)]
    stok = max((s for s in stok if s), default=None)
    emtia = [x for x in bedeller if any(a in katla(x["kalem"]) for a in ("emtia", "stok", "hammadde", "mamul"))]
    for x in bedeller:
        if x["bedel"] is None:
            b("orta", f"Sigorta bedeli yazılmamış: {x['kalem']}", "bedeller")
    if stok and emtia:
        eb = sum((x["bedel"] or SIFIR for x in emtia), SIFIR)
        if eb < stok:
            b("orta", f"Emtia bedeli ({tl(eb)} TL) raporda yazan maksimum stok değerinin ({tl(stok)} TL) altında: eksik sigorta (oransal tazminat) "
                      "riski; değerleri sigortalıyla teyit edin", "bedeller / rapor")
    elif stok and bedeller and not emtia:
        b("orta", f"Raporda maksimum stok değeri {tl(stok)} TL yazıyor ama bedeller tablosunda emtia kalemi yok", "bedeller / rapor")
    if not bedeller:
        b("bilgi", "Sigorta bedelleri tablosu verilmedi", "")
    return {"bulgular": bulgular, "liste": liste, "adaylar": adaylar, "oneriler": oneriler, "bedeller": bedeller, "bedel": bedel, "stok": stok}


# ----------------------------------------------------------------------------
# Maskeleme
# ----------------------------------------------------------------------------

def takma_adlar(raporlar: list[Dosya], terimler: list[str]) -> dict[str, str]:
    harita = {}
    tum = "\n".join(d.metin for d in raporlar)
    sig = anahtarli(tum, "Sigortalı", "Sigorta Ettiren", "Firma")
    if sig:
        harita[sig] = "[SİGORTALI]"
        kok = re.split(r"\s+(?:San\.|Sanayi|Tic\.|Ticaret|A\.Ş\.|Ltd\.)", sig)[0].strip()
        if len(kok) > 4 and kok != sig:
            harita[kok] = "[SİGORTALI]"
    adres = anahtarli(tum, "Risk Adresi", "Adres")
    if adres:
        harita[adres] = "[ADRES]"
    kisiler = [re.split(r"\s*[(,–-]", x)[0].strip() for x in [anahtarli(tum, "Görüşülen", "Görüşülen Kişi", "Yetkili")] if x]
    for i, ad in enumerate(dict.fromkeys(a.strip() for a in kisiler + list(terimler) if a.strip()), 1):
        harita.setdefault(ad, f"[KİŞİ-{i}]")
    return harita


def maskele(metin: str, harita: dict[str, str]) -> str:
    for gercek, takma in sorted(harita.items(), key=lambda x: -len(x[0])):
        metin = re.sub(re.escape(gercek), takma, metin, flags=re.I)
    return llm.maskele(metin)


def geri_ac(metin: str, harita: dict[str, str]) -> str:
    for gercek, takma in sorted(harita.items(), key=lambda x: (-len(x[1]), -len(x[0]))):
        metin = metin.replace(takma, gercek)
    return metin


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

ALINTILI = {"type": "object", "properties": {"aciklama": {"type": "string"}, "alinti": {"type": "string"}},
            "required": ["aciklama", "alinti"], "additionalProperties": False}
SEMA = {
    "type": "object",
    "properties": {
        "tesis_ozeti": {"type": "string"},
        "riskler": {"type": "array", "items": {
            "type": "object",
            "properties": {"tur": {"type": "string", "enum": RISK_TURLERI},
                           "seviye": {"type": "string", "enum": ["Yüksek", "Orta", "Düşük", "Bilgi yok"]},
                           "olumlu": {"type": "array", "items": ALINTILI}, "olumsuz": {"type": "array", "items": ALINTILI},
                           "degerlendirme": {"type": "string"}},
            "required": ["tur", "seviye", "olumlu", "olumsuz", "degerlendirme"], "additionalProperties": False}},
        "iyilestirmeler": {"type": "array", "items": {
            "type": "object",
            "properties": {"oneri": {"type": "string"}, "risk_turu": {"type": "string", "enum": RISK_TURLERI},
                           "oncelik": {"type": "string", "enum": ["Kabul öncesi", "Kısa vade", "Orta vade", "Uzun vade"]},
                           "kaynak": {"type": "string", "enum": ["Rapor", "Model önerisi"]}, "alinti": {"type": "string"}},
            "required": ["oneri", "risk_turu", "oncelik", "kaynak", "alinti"], "additionalProperties": False}},
        "kabul_onerisi": {"type": "object", "properties": {
            "karar": {"type": "string", "enum": ["Kabul", "Şartlı kabul", "Ek bilgi / yeniden inceleme", "Ret"]},
            "gerekce": {"type": "string"}, "sartlar": {"type": "array", "items": {"type": "string"}}},
            "required": ["karar", "gerekce", "sartlar"], "additionalProperties": False},
        "eksik_bilgiler": {"type": "array", "items": {"type": "string"}},
    },
    "required": ["tesis_ozeti", "riskler", "iyilestirmeler", "kabul_onerisi", "eksik_bilgiler"],
    "additionalProperties": False,
}


def sayilar(metin: str) -> set[str]:
    return {re.sub(r"\D", "", m) for m in re.findall(r"\d[\d.,]*\d", metin) if len(re.sub(r"\D", "", m)) >= 3}


def kod_ozeti(k: dict) -> list[str]:
    satir = [f"- Kontrol listesi: {sum(x['var'] for x in k['liste'])}/{len(k['liste'])} konu raporda var"]
    satir += [f"- Raporda bilgi yok: {x['tur']} — {x['konu']}" for x in k["liste"] if not x["var"]]
    satir += [f"- Bedel: {x['kalem']} {tl(x['bedel'])} TL" for x in k["bedeller"]]
    if k["bedeller"]:
        satir.append(f"- Toplam sigorta bedeli: {tl(k['bedel']['toplam'])} TL")
    if k["stok"]:
        satir.append(f"- Raporda maksimum stok değeri: {tl(k['stok'])} TL")
    satir.append(f"- Rapordaki numaralı öneri sayısı: {len(k['oneriler'])}")
    return satir


def hazirla(raporlar: list[Dosya], k: dict, harita: dict) -> tuple[str, str]:
    mesaj = "\n".join([
        "<rapor>", *[f'<belge ad="{d.ad}">\n{maskele(d.metin, harita)}\n</belge>' for d in raporlar], "</rapor>",
        "<kod_ozeti>", *[maskele(x, harita) for x in kod_ozeti(k)], "</kod_ozeti>",
        "<kod_kontrolleri>", *([f"- [{o}] {maskele(a, harita)}" for o, a, _ in k["bulgular"]] or ["- Kontrol bulgusu yok"]), "</kod_kontrolleri>",
    ])
    return (BURASI / "prompt.md").read_text(encoding="utf-8"), mesaj


def denetle(y: dict, k: dict, rapor_metni: str, mesaj: str) -> list[tuple[str, str, str]]:
    sorunlar = []
    kaynak = duz(rapor_metni)

    def dogrula(alinti: str) -> bool:
        return bool(alinti.strip()) and duz(alinti) in kaynak

    for r in y["riskler"]:
        for taraf in ("olumlu", "olumsuz"):
            for x in r[taraf]:
                x["dogrulandi"] = dogrula(x["alinti"])
                if not x["dogrulandi"]:
                    sorunlar.append(("yüksek", f"{r['tur']} ({taraf}): alıntı raporda birebir geçmiyor — \"{x['alinti'][:90]}\"", "model ↔ rapor"))
    for x in y["iyilestirmeler"]:
        x["dogrulandi"] = dogrula(x["alinti"]) if x["kaynak"] == "Rapor" else None
        if x["kaynak"] == "Rapor" and not x["dogrulandi"]:
            sorunlar.append(("yüksek", f"Rapor kaynaklı gösterilen öneri raporda bulunamadı: {x['oneri'][:100]}", "model ↔ rapor"))
    alintilar = [duz(x["alinti"]) for x in y["iyilestirmeler"] if x["alinti"].strip()]
    for o in k["oneriler"]:
        do = duz(o)
        if not any(a and (a in do or do in a or (len(a) > 20 and a[:40] in do)) for a in alintilar):
            sorunlar.append(("orta", f"Rapordaki öneri özette yok: {o[:120]}", "rapor ↔ model"))
    karar = y["kabul_onerisi"]["karar"]
    yuksek = [r["tur"] for r in y["riskler"] if r["seviye"] == "Yüksek"]
    oncesi = [x["oneri"] for x in y["iyilestirmeler"] if x["oncelik"] == "Kabul öncesi"]
    if karar == "Kabul" and (yuksek or oncesi):
        sorunlar.append(("yüksek", f"Öneri 'Kabul' ama " + (f"yüksek riskler var ({', '.join(yuksek)})" if yuksek else f"{len(oncesi)} kabul öncesi iyileştirme var"),
                         "model"))
    if karar == "Şartlı kabul" and not y["kabul_onerisi"]["sartlar"]:
        sorunlar.append(("orta", "Öneri 'Şartlı kabul' ama şart yazılmamış", "model"))
    bilinen = sayilar(mesaj) | {x for x in re.findall(r"\d+", mesaj) if len(x) >= 3}
    metin = " ".join([y["tesis_ozeti"], y["kabul_onerisi"]["gerekce"], *y["kabul_onerisi"]["sartlar"], *y["eksik_bilgiler"],
                      *[r["degerlendirme"] for r in y["riskler"]], *[x["aciklama"] for r in y["riskler"] for t in ("olumlu", "olumsuz") for x in r[t]],
                      *[x["oneri"] for x in y["iyilestirmeler"]]])
    y["dogrulanamayan_sayilar"] = sorted(s for s in sayilar(metin) if s not in bilinen)
    return sorunlar


def model_yaz(raporlar, k, harita) -> dict:
    sistem, mesaj = hazirla(raporlar, k, harita)
    y = llm.json_iste(sistem, mesaj, SEMA)
    y["sorunlar"] = denetle(y, k, "\n".join(maskele(d.metin, harita) for d in raporlar), mesaj)
    return y


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"yüksek": "FDE2E1", "orta": "FFF4CE", "bilgi": "E8F0FE", "Yüksek": "FDE2E1", "Orta": "FFF4CE", "Düşük": "E3F4E1", "Bilgi yok": "EEEEEE"}
KARAR = PatternFill("solid", fgColor="FFF4CE")
MODEL = PatternFill("solid", fgColor="EDE7FF")
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def rapor_yaz(cikti: Path, raporlar, atlanan, k: dict, y: dict, harita: dict) -> Path:
    g = lambda s: geri_ac(s, harita)  # noqa: E731
    tum = k["bulgular"] + y["sorunlar"]
    sig = next((a for a, t in harita.items() if t == "[SİGORTALI]"), "—")
    ko = y["kabul_onerisi"]
    wb = Workbook()
    o = wb.active
    o.title = "Özet"
    o.column_dimensions["A"].width, o.column_dimensions["B"].width = 30, 100
    for a, d in [("Sigortalı", sig), ("Kaynaklar", "; ".join(d.ad for d in raporlar)), ("Okunamayan", "; ".join(atlanan) or "—"),
                 ("Toplam sigorta bedeli", f"{tl(k['bedel']['toplam'])} TL" if k["bedeller"] else "—"),
                 ("Tesis özeti (model)", g(y["tesis_ozeti"])),
                 ("Risk seviyeleri (model)", " · ".join(f"{r['tur']}: {r['seviye']}" for r in y["riskler"])),
                 ("Kabul önerisi (taslak, model)", f"{ko['karar']} — {g(ko['gerekce'])}"),
                 ("Şartlar (model)", "\n".join(f"- {g(s)}" for s in ko["sartlar"]) or "—"),
                 ("Eksik bilgiler (model)", "\n".join(f"- {g(s)}" for s in y["eksik_bilgiler"]) or "—"),
                 ("Doğrulanamayan sayılar", ", ".join(y["dogrulanamayan_sayilar"]) or "—"), ("Model", llm.kullanim_ozeti()),
                 ("Underwriter Kararı", ""), ("Karar Notu", "")]:
        o.append([a, d])
        o.cell(o.max_row, 1).font = Font(bold=True)
        if "model" in a:
            o.cell(o.max_row, 2).fill = MODEL
        if a.startswith(("Underwriter", "Karar")):
            o.cell(o.max_row, 2).fill = KARAR
        if a == "Doğrulanamayan sayılar" and y["dogrulanamayan_sayilar"]:
            o.cell(o.max_row, 2).fill = PatternFill("solid", fgColor=RENK["yüksek"])
        for h in o[o.max_row]:
            h.alignment = UST

    rs = wb.create_sheet("Riskler")
    _baslik(rs, ["Risk", "Seviye (taslak)", "Yön", "Gözlem", "Rapordan Alıntı", "Doğrulandı", "Değerlendirme"], (16, 12, 9, 44, 60, 11, 50))
    for r in y["riskler"]:
        satirlar = [("Olumlu", x) for x in r["olumlu"]] + [("Olumsuz", x) for x in r["olumsuz"]] or [("", {"aciklama": "", "alinti": "", "dogrulandi": None})]
        for i, (yon, x) in enumerate(satirlar):
            rs.append([r["tur"], r["seviye"], yon, g(x["aciklama"]), g(x["alinti"]), {True: "Evet", False: "HAYIR", None: ""}[x["dogrulandi"]],
                       g(r["degerlendirme"]) if i == 0 else ""])
            rs.cell(rs.max_row, 2).fill = PatternFill("solid", fgColor=RENK[r["seviye"]])
            if x["dogrulandi"] is False:
                rs.cell(rs.max_row, 6).fill = PatternFill("solid", fgColor=RENK["yüksek"])
            for h in rs[rs.max_row]:
                h.alignment = UST

    iy = wb.create_sheet("İyileştirmeler")
    _baslik(iy, ["No", "Öneri", "Risk", "Öncelik", "Kaynak", "Rapordan Alıntı", "Doğrulandı", "Takip / Sigortalı Cevabı"], (5, 60, 16, 14, 14, 60, 11, 30))
    sira = {"Kabul öncesi": 0, "Kısa vade": 1, "Orta vade": 2, "Uzun vade": 3}
    for n, x in enumerate(sorted(y["iyilestirmeler"], key=lambda x: sira[x["oncelik"]]), 1):
        iy.append([n, g(x["oneri"]), x["risk_turu"], x["oncelik"], x["kaynak"], g(x["alinti"]), {True: "Evet", False: "HAYIR", None: "—"}[x["dogrulandi"]], ""])
        if x["oncelik"] == "Kabul öncesi":
            iy.cell(iy.max_row, 4).fill = PatternFill("solid", fgColor=RENK["orta"])
        if x["kaynak"] == "Model önerisi":
            iy.cell(iy.max_row, 5).fill = MODEL
        iy.cell(iy.max_row, 8).fill = KARAR
        for h in iy[iy.max_row]:
            h.alignment = UST

    kl = wb.create_sheet("Kontrol Listesi")
    _baslik(kl, ["Risk", "Konu", "Raporda Var mı?", "İlk Geçtiği Cümle"], (16, 36, 14, 90))
    for x in k["liste"]:
        kl.append([x["tur"], x["konu"], "Evet" if x["var"] else "Hayır", g(x["cumle"])])
        if not x["var"]:
            kl.cell(kl.max_row, 3).fill = PatternFill("solid", fgColor=RENK["orta"])

    kt = wb.create_sheet("Kontroller")
    _baslik(kt, ["Önem", "Kontrol", "Kaynak", "İnceleme"], (9, 100, 18, 26))
    for on, a, kay in tum:
        kt.append([on, a, kay, ""])
        kt.cell(kt.max_row, 1).fill = PatternFill("solid", fgColor=RENK[on])
        kt.cell(kt.max_row, 4).fill = KARAR
        kt.cell(kt.max_row, 2).alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

    md = cikti.with_suffix(".md")
    sat = [f"# Risk Değerlendirme Özeti (taslak) — {sig}", "", f"Kaynak: {', '.join(d.ad for d in raporlar)}", "", "## Tesis", "", g(y["tesis_ozeti"]), ""]
    if k["bedeller"]:
        sat += ["| Kalem | Bedel (TL) |", "|---|---:|"] + [f"| {x['kalem']} | {tl(x['bedel'])} |" for x in k["bedeller"]]
        sat += [f"| **Toplam** | **{tl(k['bedel']['toplam'])}** |", ""]
    sat += ["## Riskler", "", "| Risk | Seviye (taslak) |", "|---|---|"] + [f"| {r['tur']} | {r['seviye']} |" for r in y["riskler"]] + [""]
    for r in y["riskler"]:
        sat += [f"### {r['tur']} — {r['seviye']}", "", g(r["degerlendirme"]), ""]
        sat += [f"- (+) {g(x['aciklama'])} — \"{g(x['alinti'])}\"" + ("" if x["dogrulandi"] else " **[alıntı doğrulanamadı]**") for x in r["olumlu"]]
        sat += [f"- (−) {g(x['aciklama'])} — \"{g(x['alinti'])}\"" + ("" if x["dogrulandi"] else " **[alıntı doğrulanamadı]**") for x in r["olumsuz"]]
        sat += [""]
    sat += ["## İyileştirmeler", "", "| Öncelik | Öneri | Risk | Kaynak |", "|---|---|---|---|"]
    sat += [f"| {x['oncelik']} | {g(x['oneri'])} | {x['risk_turu']} | {x['kaynak']} |" for x in sorted(y["iyilestirmeler"], key=lambda x: sira[x["oncelik"]])]
    sat += ["", "## Kabul önerisi (taslak)", "", f"**{ko['karar']}** — {g(ko['gerekce'])}", ""] + [f"- {g(s)}" for s in ko["sartlar"]]
    if y["eksik_bilgiler"]:
        sat += ["", "## Eksik bilgiler", ""] + [f"- {g(s)}" for s in y["eksik_bilgiler"]]
    sat += ["", "## Kod kontrolleri", ""] + ([f"- **{o_}** {g(a)}" for o_, a, _ in tum] or ["- Bulgu yok"])
    if y["dogrulanamayan_sayilar"]:
        sat += ["", f"> Girdilerde bulunmayan sayılar (kontrol edin): {', '.join(y['dogrulanamayan_sayilar'])}"]
    sat += ["", "> Taslaktır; risk kabul kararı yetkili underwriter'a aittir. Tarife, teminat ve muafiyetleri şirketinizin "
                "risk kabul kurallarına göre belirleyin."]
    md.write_text("\n".join(sat) + "\n", encoding="utf-8")
    return md


def calistir(klasor: Path, cikti: Path, terimler: list[str] | None = None, evet: bool = False) -> dict:
    if not klasor.is_dir():
        raise llm.LLMHatasi(f"{klasor} klasörü bulunamadı.")
    raporlar, bedeller, atlanan = dosyalari_oku(klasor)
    if not raporlar:
        raise llm.LLMHatasi(f"{klasor}: okunabilir rapor yok (.txt/.pdf/.docx).")
    k = kontrol_et(raporlar, bedeller)
    harita = takma_adlar(raporlar, terimler or [])
    print(f"[OK] {len(raporlar)} belge · kontrol listesi {sum(x['var'] for x in k['liste'])}/{len(k['liste'])} · {len(k['oneriler'])} rapor önerisi")
    llm.onay_al(f"{len(raporlar)} belgenin metni gönderilecek (sigortalı unvanı, risk adresi ve "
                f"{sum(t.startswith('[KİŞİ') for t in harita.values())} kişi takma adlı; telefon, e-posta, TCKN, IBAN maskeli).", evet)
    y = model_yaz(raporlar, k, harita)
    md = rapor_yaz(cikti, raporlar, atlanan, k, y, harita)
    return {"kontrol": k, "yanit": y, "harita": harita, "md": md, "raporlar": raporlar}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Risk inceleme (survey) raporundan risk değerlendirme özeti ve kabul önerisi taslağı hazırlar.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "ornek_plastik", help="Klasör: survey raporu (+ isteğe bağlı bedeller tablosu)")
    p.add_argument("--gizle", nargs="*", default=[], metavar="AD", help="Ayrıca maskelenecek adlar")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "risk_degerlendirme_ozeti.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.cikti, a.gizle, a.evet)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    for o, ac, _ in s["kontrol"]["bulgular"] + s["yanit"]["sorunlar"]:
        if o == "yüksek":
            print(f"[X] {ac}")
    if s["yanit"]["dogrulanamayan_sayilar"]:
        print(f"[!] Girdilerde olmayan sayılar: {', '.join(s['yanit']['dogrulanamayan_sayilar'])}")
    print(f"[i] Kabul önerisi (taslak): {s['yanit']['kabul_onerisi']['karar']}")
    print(f"[OK] Rapor: {a.cikti.resolve()} · {s['md'].resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
