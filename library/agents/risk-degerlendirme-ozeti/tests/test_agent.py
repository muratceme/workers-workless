"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

DOSYA = KLASOR / "ornek_veri" / "ornek_plastik"
MESAJLAR = []
R = {i: x for i, x in enumerate([
    "R1. Üretim holüne otomatik sprinkler sistemi kurulmalıdır (orta vade).",
    "R2. Elektrik panolarında termal kamera ölçümü yaptırılmalı ve tespit edilen sıcak noktalar giderilmelidir (kısa vade).",
    "R3. Bitmiş ürün paletleri bina duvarından en az 10 m uzağa taşınmalıdır (poliçe öncesi).",
    "R4. Sıcak çalışmalar için yazılı izin prosedürü uygulanmalıdır (kısa vade).",
    "R5. Eksik yangın dolabı hortumları tamamlanmalıdır (poliçe öncesi).",
    "R6. Depo raf sistemleri zemine ankrajlanmalıdır (orta vade)."], 1)}   # R7 bilerek yok


def iy(no, oncelik, risk="Yangın"):
    return {"oneri": R[no][4:], "risk_turu": risk, "oncelik": oncelik, "kaynak": "Rapor", "alinti": R[no]}


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return {
        "tesis_ozeti": "[SİGORTALI] polietilen film üretiyor; 6.000 m² çelik konstrüksiyon üretim holü, 85 çalışan.",
        "riskler": [
            {"tur": "Yangın", "seviye": "Yüksek",
             "olumlu": [{"aciklama": "Yangın suyu deposu var.", "alinti": "Yangın suyu için 200 m³ kapasiteli depo ve dizel yangın pompası mevcuttur."}],
             "olumsuz": [{"aciklama": "Üretim holünde sprinkler yok.", "alinti": "üretim holünde bulunmamaktadır"},
                         {"aciklama": "Yanıcı panel.", "alinti": "Çatı ve cepheler yanıcı sandviç panel ile kaplıdır."}],
             "degerlendirme": "Yanıcı yapı ve sprinklersiz üretim holü yangın riskini artırıyor; 2023'te 250.000 TL hasar oldu."},
            {"tur": "Deprem", "seviye": "Orta", "olumlu": [], "olumsuz": [{"aciklama": "Raf ankrajı yok.", "alinti": "raf sistemleri zemine ankrajlı değildir"}],
             "degerlendirme": "Zemin etüdü görülmedi."},
        ],
        "iyilestirmeler": [iy(1, "Orta vade"), iy(2, "Kısa vade"), iy(3, "Kabul öncesi"), iy(4, "Kısa vade"), iy(5, "Kabul öncesi"),
                           iy(6, "Orta vade", "Deprem"),
                           {"oneri": "Paratoner tesisatı kontrol edilmeli.", "risk_turu": "Yangın", "oncelik": "Kısa vade", "kaynak": "Model önerisi",
                            "alinti": "Paratoner tesisatı mevcuttur."}],
        "kabul_onerisi": {"karar": "Kabul", "gerekce": "Önlemler yeterli.", "sartlar": []},
        "eksik_bilgiler": ["Yıldırımdan korunma tesisatı bilgisi."],
    }


class KontrolTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.r, cls.b, cls.a = agent.dosyalari_oku(DOSYA)
        cls.k = agent.kontrol_et(cls.r, cls.b)
        cls.metin = "\n".join(f"{o} {a}" for o, a, _ in cls.k["bulgular"])

    def test_bedeller(self):
        self.assertEqual(len(self.b), 4)
        self.assertEqual(self.k["bedel"]["toplam"], D(251000000))
        self.assertEqual(self.k["stok"], D(55000000))
        self.assertIn("orta Emtia bedeli (40.000.000 TL) raporda yazan maksimum stok değerinin (55.000.000 TL) altında", self.metin)

    def test_kontrol_listesi(self):
        eksik = [x["konu"] for x in self.k["liste"] if not x["var"]]
        self.assertEqual(eksik, ["Yıldırımdan korunma"])
        self.assertIn("orta Raporda bilgi yok: Yangın — Yıldırımdan korunma", self.metin)

    def test_olumsuz_gozlemler(self):
        self.assertEqual(len(self.k["adaylar"]), 9)
        for parca in ("yangın duvarı bulunmamaktadır", "üretim holünde bulunmamaktadır", "hortumu eksiktir", "termal kamera ölçümü yapılmamıştır",
                      "izin sistemi bulunmamaktadır", "zemin etüdü raporu görülmemiştir", "ankrajlı değildir", "depolarda bulunmamaktadır"):
            self.assertIn(parca, self.metin)
        self.assertNotIn("Bodrum", self.metin)                                    # tehlikenin yokluğu olumsuz değil
        self.assertNotIn("Olumsuz gözlem: Eksik yangın dolabı", self.metin)      # öneri satırı
        self.assertEqual(len(self.k["oneriler"]), 7)

    def test_takma_adlar(self):
        h = agent.takma_adlar(self.r, ["Ek Kişi"])
        self.assertEqual((h["Örnek Plastik Ambalaj Sanayi ve Ticaret A.Ş."], h["Kurgu OSB 4. Cadde No: 12, Örnekköy"], h["Murat Denemeci"], h["Ek Kişi"]),
                         ("[SİGORTALI]", "[ADRES]", "[KİŞİ-1]", "[KİŞİ-2]"))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_uctan_uca(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "r.xlsx"
            s = agent.calistir(DOSYA, cikti, evet=True)
            m = MESAJLAR[0]
            for gizli in ("Örnek Plastik", "Kurgu OSB", "Murat Denemeci"):
                self.assertNotIn(gizli, m)
            for var in ("[SİGORTALI]", "[ADRES]", "[KİŞİ-1]", "Raporda bilgi yok: Yangın — Yıldırımdan korunma"):
                self.assertIn(var, m)
            y = s["yanit"]
            yan = y["riskler"][0]
            self.assertEqual([x["dogrulandi"] for x in yan["olumlu"] + yan["olumsuz"]], [True, True, False])
            self.assertTrue(y["riskler"][1]["olumsuz"][0]["dogrulandi"])
            self.assertEqual([x["dogrulandi"] for x in y["iyilestirmeler"]], [True] * 6 + [None])
            sorun = " ".join(a for _, a, _ in y["sorunlar"])
            self.assertIn("Yangın (olumsuz): alıntı raporda birebir geçmiyor", sorun)
            self.assertIn("Rapordaki öneri özette yok: R7. Hammadde deposu", sorun)
            self.assertNotIn("R6.", sorun)
            self.assertIn("Öneri 'Kabul' ama yüksek riskler var (Yangın)", sorun)
            self.assertEqual(y["dogrulanamayan_sayilar"], ["250000"])
            md = s["md"].read_text(encoding="utf-8")
            self.assertIn("Örnek Plastik Ambalaj Sanayi ve Ticaret A.Ş. polietilen film üretiyor", md)
            self.assertIn("**[alıntı doğrulanamadı]**", md)
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Özet", "Riskler", "İyileştirmeler", "Kontrol Listesi", "Kontroller"])
            self.assertEqual(wb["İyileştirmeler"]["D2"].value, "Kabul öncesi")
            self.assertEqual(wb["İyileştirmeler"]["H1"].value, "Takip / Sigortalı Cevabı")

    def test_klasor_yok(self):
        with self.assertRaises(llm.LLMHatasi):
            agent.calistir(Path("olmayan_klasor"), Path("x.xlsx"), evet=True)


if __name__ == "__main__":
    unittest.main()
