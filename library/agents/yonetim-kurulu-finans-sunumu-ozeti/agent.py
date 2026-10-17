"""
Yönetim Kurulu Finans Sunumu Özeti — Gerisi Bende AI Agent
Finans › Finans Müdürü

1. Kod dönem sonuç tablosunu (Kalem; Cari Dönem; Önceki Dönem; Bütçe) okur ve hesaplar:
   - Her kalemin önceki döneme göre değişimi (tutar ve %) ve bütçe sapması (tutar ve %)
   - Tanınan kalemlerden oranlar: brüt kâr, FAVÖK ve net kâr marjı (puan değişimiyle), cari oran
   - Bütçeden en çok sapan kalemler
   Yönetim notları (isteğe bağlı) satır satır numaralanır (N1, N2 …).
2. Model yönetim kurulu için öne çıkanlar ve bölümler (gelir, kârlılık, bütçe, nakit / borç, riskler) yazar; her madde
   dayandığı kalem, oran ya da not numarasını gösterir.
3. Kod denetler:
   - Metindeki her sayı, maddenin dayandığı kalem / oran / notun hazır rakamlarından biri mi (yuvarlama ya da yeni hesap yok)
   - "Arttı / azaldı" yönü değişimin işaretiyle tutarlı mı
   - Neden-sonuç ifadesi ("nedeniyle", "kaynaklı") bir yönetim notuna dayanıyor mu
   - Bütçeden en çok sapan kalemler metinde ele alınmış mı; "kesinlikle", "garanti" gibi ifade var mı
   Sorun varsa modele geri gönderilir (en fazla --tur kez).

Kullanım:
    python agent.py                                                     # örnek çeyrek sonuçları
    python agent.py --girdi sonuclar.xlsx --notlar notlar.txt --donem "2026 3. Çeyrek" --birim "milyon TL"
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
REF = {"type": "array", "items": S}
_M = {"type": "object", "properties": {"metin": S, "dayanak": REF}, "required": ["metin", "dayanak"], "additionalProperties": False}
SEMA = {
    "type": "object",
    "properties": {
        "one_cikanlar": {"type": "array", "items": _M},
        "bolumler": {"type": "array", "items": {"type": "object", "properties": {
            "baslik": {"type": "string", "enum": ["Gelir", "Kârlılık", "Bütçe Gerçekleşme", "Nakit ve Borç", "Riskler ve Dikkat Noktaları"]},
            "maddeler": {"type": "array", "items": _M}}, "required": ["baslik", "maddeler"], "additionalProperties": False}},
        "karar_gerektiren_konular": {"type": "array", "items": _M},
    },
    "required": ["one_cikanlar", "bolumler", "karar_gerektiren_konular"],
    "additionalProperties": False,
}
TANIMA = {"hasilat": ("hasilat", "net satis", "satis gelir", "gelirler"), "brut": ("brut kar",), "favok": ("favok", "ebitda"),
          "net": ("net donem kar", "net kar", "donem net kar"), "donen": ("donen varlik",), "kvy": ("kisa vadeli yukumluluk", "kisa vadeli borc")}
ARTIS = re.compile(r"\b(art(?:tı|ış|arak|mış)\w*|yüksel\w*|büyü\w*|iyileş\w*)", re.I)
AZALIS = re.compile(r"\b(azal\w*|geril\w*|düş(?:tü|üş|erek|müş)\w*|daral\w*|kötüleş\w*)", re.I)
NEDEN = re.compile(r"\b(nedeniyle|sebebiyle|kaynaklı|kaynaklanan|etkisiyle|dolayı|sayesinde|bağlı olarak)\b", re.I)
KESIN = re.compile(r"\b(kesinlikle|garanti\w*|mutlaka|şüphesiz)\b", re.I)
EN_COK = 3


def sayi(v) -> float | None:
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v or "").strip().replace(" ", "")
    if not s or s in {"-", "—"}:
        return None
    neg = s.startswith("(") and s.endswith(")") or s.startswith("-")
    s = s.strip("()-")
    try:
        x = float(sayi_norm(s))
    except ValueError:
        return None
    return -x if neg else x


def tr(v: float | None, ondalik: int = 1) -> str:
    if v is None:
        return "a.d."
    if ondalik == 1 and float(v).is_integer():
        ondalik = 0
    return f"{abs(v):,.{ondalik}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def isaret(v: float | None) -> str:
    return "" if v is None else ("+" if v > 0 else ("−" if v < 0 else "±"))


def ts(v: float | None, ondalik: int = 1) -> str:
    """İşaretli tutar: −905; artı işareti yazılmaz."""
    return ("−" if v is not None and v < 0 else "") + tr(v, ondalik)


def degisim(v: float | None, yuzde: float | None) -> str:
    return f"{isaret(v)}{tr(v)}" + ("" if yuzde is None else f" ({isaret(yuzde)}%{tr(yuzde)})")


def tabloyu_oku(yol: Path) -> list[dict]:
    r = kayitlar(yol, {"kalem": ("Kalem", "Hesap", "Gösterge"), "cari": ("Cari Dönem", "Cari", "Gerçekleşen"), "onceki": ("Önceki Dönem", "Önceki", "Geçen Yıl"),
                       "butce": ("Bütçe", "Hedef")}, ("kalem", "cari"))
    out = []
    for x in r:
        if not str(x["kalem"]).strip() or sayi(x["cari"]) is None:
            continue
        k = {"id": f"K{len(out) + 1}", "kalem": str(x["kalem"]).strip(), "cari": sayi(x["cari"]), "onceki": sayi(x["onceki"]), "butce": sayi(x["butce"])}
        # Gider kalemleri tabloda eksi / parantezli yazılır; değişim ve sapma büyüklük üzerinden hesaplanır (gider arttı = tutar büyüdü)
        k["gider"] = k["cari"] < 0 and all(k[a] is None or k[a] <= 0 for a in ("onceki", "butce"))
        if k["gider"]:
            k.update({a: abs(k[a]) for a in ("cari", "onceki", "butce") if k[a] is not None})
        k["degisim"] = None if k["onceki"] is None else k["cari"] - k["onceki"]
        k["degisim_yuzde"] = None if not k["onceki"] or k["onceki"] < 0 else k["degisim"] / k["onceki"] * 100
        k["sapma"] = None if k["butce"] is None else k["cari"] - k["butce"]
        k["sapma_yuzde"] = None if not k["butce"] or k["butce"] < 0 else k["sapma"] / k["butce"] * 100
        out.append(k)
    return out


def oranlar(kalemler: list[dict]) -> list[dict]:
    bul = {a: next((k for k in kalemler if any(katla(k["kalem"]).startswith(e) for e in es)), None) for a, es in TANIMA.items()}
    out = []

    def ekle(ad, pay, payda, carpan=100.0, birim="%"):
        c = pay["cari"] / payda["cari"] * carpan if payda["cari"] else None
        o = pay["onceki"] / payda["onceki"] * carpan if payda["onceki"] and pay["onceki"] is not None else None
        out.append({"id": f"R{len(out) + 1}", "oran": ad, "birim": birim, "cari": c, "onceki": o, "fark": None if c is None or o is None else c - o,
                    "dayanak": f"{pay['id']} / {payda['id']}"})

    h = bul["hasilat"]
    if h:
        for a, ad in (("brut", "Brüt kâr marjı"), ("favok", "FAVÖK marjı"), ("net", "Net kâr marjı")):
            if bul[a]:
                ekle(ad, bul[a], h)
    if bul["donen"] and bul["kvy"]:
        ekle("Cari oran", bul["donen"], bul["kvy"], 1.0, "kat")
    return out


def notlari_oku(yol: Path | None) -> list[str]:
    if not yol:
        return []
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: yönetim notları bulunamadı.")
    return [s.strip() for s in belge.metin_oku(yol).splitlines() if s.strip()]


def rakamlar(kalemler, oran, notlar) -> dict[str, set[str]]:
    """Her dayanak için metinde kullanılabilecek sayılar (biçimlendirilmiş hâlleriyle)."""
    r = {}
    for k in kalemler:
        r[k["id"]] = set().union(*(sayilar(tr(k[a])) for a in ("cari", "onceki", "butce", "degisim", "sapma") if k[a] is not None),
                                 *(sayilar(tr(k[a])) for a in ("degisim_yuzde", "sapma_yuzde") if k[a] is not None))
    for o in oran:
        r[o["id"]] = set().union(*(sayilar(tr(o[a], 2 if o["birim"] == "kat" else 1)) for a in ("cari", "onceki", "fark") if o[a] is not None))
    for i, n in enumerate(notlar, 1):
        r[f"N{i}"] = sayilar(n)
    return r


def en_cok_sapan(kalemler: list[dict]) -> list[str]:
    s = sorted((k for k in kalemler if k["sapma_yuzde"] is not None), key=lambda k: -abs(k["sapma_yuzde"]))
    return [k["id"] for k in s[:EN_COK] if abs(k["sapma_yuzde"]) >= 1]


def maddeler(t: dict) -> list[tuple[str, dict]]:
    out = [("Öne çıkan", m) for m in t["one_cikanlar"]]
    for b in t["bolumler"]:
        out += [(b["baslik"], m) for m in b["maddeler"]]
    return out + [("Karar", m) for m in t["karar_gerektiren_konular"]]


def denetle(t: dict, kalemler, oran, notlar, donem: str) -> list[str]:
    sorun, r = [], rakamlar(kalemler, oran, notlar)
    kdict = {k["id"]: k for k in kalemler}
    kullanilan = set()
    for i, (yer, m) in enumerate(maddeler(t), 1):
        ad = f"M{i} ({yer})"
        gec = [d for d in m["dayanak"] if d not in r]
        if gec:
            sorun.append(f"{ad}: olmayan dayanak {gec}")
        if not m["dayanak"]:
            sorun.append(f"{ad}: dayanak (kalem / oran / not) yok")
        kullanilan |= set(m["dayanak"])
        izinli = set().union(set(), *(r.get(d, set()) for d in m["dayanak"])) | sayilar(donem)
        fazla = sorted(sayilar(m["metin"]) - izinli)
        if fazla:
            sorun.append(f"{ad}: dayandığı kalem / oran / notta olmayan sayı: {', '.join(fazla)} (hazır rakamları aynen kullanın)")
        ks = [kdict[d] for d in m["dayanak"] if d in kdict]
        if len(ks) == 1 and ks[0]["degisim"]:
            art, az = bool(ARTIS.search(m["metin"])), bool(AZALIS.search(m["metin"]))
            if art != az and (art and ks[0]["degisim"] < 0 or az and ks[0]["degisim"] > 0):
                sorun.append(f"{ad}: '{ks[0]['kalem']}' önceki döneme göre {'azaldı' if ks[0]['degisim'] < 0 else 'arttı'}; metindeki yön ters")
        if NEDEN.search(m["metin"]) and not any(d.startswith("N") for d in m["dayanak"]):
            sorun.append(f"{ad}: neden-sonuç ifadesi ('{NEDEN.search(m['metin']).group(0)}') bir yönetim notuna (N…) dayanmıyor")
        if KESIN.search(m["metin"]):
            sorun.append(f"{ad}: kesinlik / garanti ifadesi ('{KESIN.search(m['metin']).group(0)}')")
    eksik = [f"{d} {kdict[d]['kalem']}" for d in en_cok_sapan(kalemler) if d not in kullanilan]
    if eksik:
        sorun.append(f"Genel: bütçeden en çok sapan kalem ele alınmamış: {', '.join(eksik)}")
    if not t["one_cikanlar"]:
        sorun.append("Genel: öne çıkanlar yok")
    return sorun


def uret(kalemler, oran, notlar, donem, birim, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    k = "\n".join(f"{x['id']} | {x['kalem']}{' (gider; tutarlar mutlak değer)' if x['gider'] else ''} | cari {ts(x['cari'])} | önceki {ts(x['onceki'])} | "
                  f"değişim {degisim(x['degisim'], x['degisim_yuzde'])} | bütçe {ts(x['butce'])} | bütçe sapması {degisim(x['sapma'], x['sapma_yuzde'])}"
                  for x in kalemler)

    def oran_satiri(x):
        d, b = (2, " kat") if x["birim"] == "kat" else (1, "%")
        return (f"{x['id']} | {x['oran']} ({x['dayanak']}) | cari {ts(x['cari'], d)}{b} | önceki {ts(x['onceki'], d)}{b} | "
                f"fark {isaret(x['fark'])}{tr(x['fark'], d)}{'' if x['birim'] == 'kat' else ' puan'}")
    o = "\n".join(oran_satiri(x) for x in oran) or "(hesaplanamadı)"
    n = "\n".join(f"N{i} | {s}" for i, s in enumerate(notlar, 1)) or "(yönetim notu yok)"
    mesaj = (f"<donem>{donem}</donem>\n<birim>{birim}</birim>\n<kalemler>\n{k}\n</kalemler>\n<oranlar>\n{o}\n</oranlar>\n<yonetim_notlari>\n{n}\n"
             f"</yonetim_notlari>\n<en_cok_sapan>{', '.join(en_cok_sapan(kalemler)) or '(bütçe yok)'}</en_cok_sapan>" + duzeltme_blogu(duzeltme, onceki))
    return llm.json_iste(sistem, mesaj, SEMA)


def rapor_yaz(cikti: Path, t, kalemler, oran, notlar, donem, birim, kalan, tur) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    baslik(ws, ["Bölüm", "Madde", "Dayanak"], (22, 100, 16))
    for yer, m in maddeler(t):
        ws.append([yer, m["metin"], ", ".join(m["dayanak"])])
        boya(ws, [2], AI_DOLGU)
    son_not(ws, "Önemli: Rakamlar kod tarafından hesaplanmıştır; yorumlar AI taslağıdır. Yönetim kuruluna sunmadan önce finans müdürü "
                "onaylamalıdır.", kalan, tur)
    hizala(ws)
    r = wb.create_sheet("Rakamlar")
    baslik(r, ["No", "Kalem", "Cari", "Önceki", "Değişim", "Değişim %", "Bütçe", "Sapma", "Sapma %"], (5, 34, 14, 14, 14, 10, 14, 14, 10))
    sapan = set(en_cok_sapan(kalemler))
    for k in kalemler:
        r.append([k["id"], k["kalem"], k["cari"], k["onceki"], k["degisim"], None if k["degisim_yuzde"] is None else round(k["degisim_yuzde"], 1),
                  k["butce"], k["sapma"], None if k["sapma_yuzde"] is None else round(k["sapma_yuzde"], 1)])
        for c in (3, 4, 5, 7, 8):
            r.cell(r.max_row, c).number_format = "#,##0.0;(#,##0.0)"
        if k["id"] in sapan:
            boya(r, [8, 9], SARI)
    r.append([])
    baslik(r, ["No", "Oran", "Cari", "Önceki", "Fark", "Dayanak"])
    for o in oran:
        r.append([o["id"], o["oran"] + (" (kat)" if o["birim"] == "kat" else " (%)"), *(None if o[a] is None else round(o[a], 2) for a in ("cari", "onceki", "fark")),
                  o["dayanak"]])
    hizala(r)
    if notlar:
        n = wb.create_sheet("Yönetim Notları")
        baslik(n, ["No", "Not"], (6, 100))
        for i, s in enumerate(notlar, 1):
            n.append([f"N{i}", s])
        hizala(n)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    metin = [f"YÖNETİM KURULU FİNANS ÖZETİ — {donem} ({birim})", "", "ÖNE ÇIKANLAR"] + [f"• {m['metin']}" for m in t["one_cikanlar"]]
    for b in t["bolumler"]:
        metin += ["", b["baslik"].upper()] + [f"• {m['metin']}" for m in b["maddeler"]]
    if t["karar_gerektiren_konular"]:
        metin += ["", "KARAR GEREKTİREN KONULAR"] + [f"• {m['metin']}" for m in t["karar_gerektiren_konular"]]
    metin += ["", "Taslaktır — finans müdürü onayı olmadan dağıtılmaz."]
    cikti.with_suffix(".txt").write_text("\n".join(metin), encoding="utf-8")
    return "\n".join(metin)


def calistir(girdi: Path, cikti: Path, donem: str, birim: str, notlar_yolu: Path | None = None, tur: int = 2, evet: bool = False) -> dict:
    kalemler = tabloyu_oku(girdi)
    if not kalemler:
        raise llm.LLMHatasi(f"{girdi.name}: sayısal kalem yok.")
    oran = oranlar(kalemler)
    notlar = notlari_oku(notlar_yolu)
    print(f"[OK] {len(kalemler)} kalem · {len(oran)} oran · {len(notlar)} yönetim notu · en çok sapan: {', '.join(en_cok_sapan(kalemler)) or '-'}")
    llm.onay_al(f"{len(kalemler)} kalemin dönem rakamları, hesaplanan oranlar ve {len(notlar)} yönetim notu gönderilecek.", evet)
    t, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(kalemler, oran, notlar, donem, birim, d, o), lambda x: denetle(x, kalemler, oran, notlar, donem), tur)
    metin = rapor_yaz(cikti, t, kalemler, oran, notlar, donem, birim, kalan, yapilan)
    return {"kalemler": kalemler, "oranlar": oran, "taslak": t, "kalan": kalan, "tur": yapilan, "metin": metin}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    o = BURASI / "ornek_veri"
    p = argparse.ArgumentParser(description="Dönem finansal sonuçlarından yönetim kurulu için kısa, yorumlu özet hazırlar.")
    p.add_argument("--girdi", type=Path, default=o / "donem_sonuclari.csv", help="Sonuç tablosu (.xlsx/.csv): Kalem, Cari Dönem; Önceki Dönem, Bütçe")
    p.add_argument("--notlar", type=Path, default=o / "yonetim_notlari.txt", help="Yönetim notları (.txt/.docx); her satır bir açıklama (isteğe bağlı)")
    p.add_argument("--donem", default="2026 3. Çeyrek", help="Dönem adı")
    p.add_argument("--birim", default="milyon TL", help="Tablodaki tutarların birimi")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "yk_finans_ozeti.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.cikti, a.donem, a.birim, a.notlar, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(maddeler(s['taslak']))} madde")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()} (+ {a.cikti.with_suffix('.txt').name})")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
