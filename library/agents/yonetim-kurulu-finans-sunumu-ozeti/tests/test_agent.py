"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

TABLO = KLASOR / "ornek_veri" / "donem_sonuclari.csv"
NOTLAR = KLASOR / "ornek_veri" / "yonetim_notlari.txt"
DONEM = "2026 3. Çeyrek"
MESAJLAR = []


def m(metin, *dayanak):
    return {"metin": metin, "dayanak": list(dayanak)}


def taslak(dogru=True) -> dict:
    one = [m("Hasılat önceki yılın aynı dönemine göre %15,7 artışla 1.250 milyon TL'ye yükseldi; artış ağırlıklı olarak ihracat satışlarındaki "
             "%22'lik büyümeden kaynaklandı.", "K1", "N1"),
           m("Net dönem kârı %17,4 azalarak 71 milyon TL oldu ve bütçenin %32,4 altında kaldı.", "K8"),
           m("Finansman giderleri, 150 milyon TL tutarındaki yeni yatırım kredisi nedeniyle %51,2 artarak 62 milyon TL'ye çıktı ve bütçeyi %24 aştı.",
             "K7", "N2")]
    bol = [{"baslik": "Kârlılık", "maddeler": [m("Brüt kâr marjı 2,0 puan daralarak %27,6 oldu; hammadde fiyat artışları satış fiyatlarına tam "
                                                  "yansıtılamadı.", "R1", "N3")]},
           {"baslik": "Bütçe Gerçekleşme", "maddeler": [m("Esas faaliyet kârı 155 milyon TL ile bütçenin %18,4 altında gerçekleşti.", "K6")]},
           {"baslik": "Nakit ve Borç", "maddeler": [m("Net borç 45 milyon TL artışla 340 milyon TL'ye yükseldi.", "K10"),
                                                     m("Cari oran 1,34 kat ile önceki dönemin (1,41 kat) altında.", "R4")]}]
    karar = [m("Ertelenen fiyat artışı için fiyat politikası kararı gerekmektedir.", "N4")]
    if not dogru:
        one[0] = m("Hasılat %16 artışla 1.250 milyon TL'ye yükseldi.", "K1")
        one[1] = m("Net dönem kârı 71 milyon TL'ye yükseldi.", "K8")
        one[2] = m("Finansman giderleri yeni kredi nedeniyle %51,2 arttı.", "K7")
        bol[1]["maddeler"] = []
        karar = [m("Fiyat artışı kesinlikle yapılmalıdır.", "N4")]
    return {"one_cikanlar": one, "bolumler": bol, "karar_gerektiren_konular": karar}


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return taslak("<duzeltme>" in kullanici)


def veri():
    k = agent.tabloyu_oku(TABLO)
    return k, agent.oranlar(k), agent.notlari_oku(NOTLAR)


class KodTesti(unittest.TestCase):
    def test_hesap(self):
        k, o, _ = veri()
        self.assertEqual(len(k), 12)
        self.assertEqual((k[1]["gider"], k[1]["cari"], k[1]["degisim"], round(k[1]["degisim_yuzde"], 1)), (True, 905.0, 145.0, 19.1))
        self.assertEqual([round(x["cari"], 2) for x in o], [27.6, 15.84, 5.68, 1.34])
        self.assertEqual(agent.en_cok_sapan(k), ["K8", "K7", "K6"])
        self.assertEqual((agent.tr(1250.0), agent.tr(-15.7), agent.tr(None)), ("1.250", "15,7", "a.d."))

    def test_denetim(self):
        k, o, n = veri()
        self.assertEqual(agent.denetle(taslak(), k, o, n, DONEM), [])
        s = agent.denetle(taslak(False), k, o, n, DONEM)
        self.assertEqual(len(s), 5, s)
        self.assertTrue(any("sayı: 16" in x for x in s) and any("Net Dönem Kârı" in x and "ters" in x for x in s)
                        and any("nedeniyle" in x for x in s) and any("K6" in x for x in s) and any("kesinlikle" in x for x in s))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "yk.xlsx"
            s = agent.calistir(TABLO, cikti, DONEM, "milyon TL", NOTLAR, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            self.assertIn("K7 | Finansman Giderleri (gider; tutarlar mutlak değer) | cari 62 | önceki 41 | değişim +21 (+%51,2)", MESAJLAR[0])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Rakamlar", "Yönetim Notları"])
            self.assertIn("KARAR GEREKTİREN KONULAR", cikti.with_suffix(".txt").read_text(encoding="utf-8"))

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(TABLO, Path(tmp) / "yk.xlsx", DONEM, "milyon TL", NOTLAR, tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 5)
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--evet", "--notlar", str(NOTLAR.with_name("yok.txt"))]), 1)


if __name__ == "__main__":
    unittest.main()
