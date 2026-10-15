"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri" / "anketler.csv"
MESAJLAR = []
HEK, HEM, BEK, RAN, BIL, TEM, YEM, PER, FIZ, OTO = agent.TEMALAR[:10]
Y = {"A-001": ([HEK, BEK], "Nötr", "1 saat bekledim"), "A-004": ([BEK, FIZ], "Olumsuz", "oturacak yer yok"), "A-007": ([HEK], "Olumlu", "her şeyi güzelce anlattı"),
     "A-013": ([HEK], "Olumsuz", "sorularıma cevap vermedi"), "A-016": ([RAN], "Olumlu", "Randevu saatinde alındım"), "A-025": ([OTO], "Olumsuz", "Otopark bulmak çok zor"),
     "A-028": ([BEK], "Olumsuz", "2 saat bekledim"), "A-031": ([BEK, BIL], "Olumsuz", "kimse bilgi vermiyor"), "A-034": ([HEM], "Olumlu", "Hemşireler çok ilgiliydi"),
     "A-037": ([TEM], "Olumsuz", "banyo kirliydi"), "A-040": ([YEM], "Olumsuz", "Yemekler soğuk geliyordu"), "A-043": ([TEM], "Olumsuz", "temizliği yetersizdi")}


def iy(birim, baslik, dayanak, soru):
    return {"birim": birim, "baslik": baslik, "gerekce": "g", "dayanak": dayanak, "dusuk_soru": soru, "oneri": "ö"}


def taslak(dogru=True) -> dict:
    t = {"yorumlar": [{"id": i, "temalar": te, "duygu": d, "alinti": a} for i, (te, d, a) in Y.items()],
         "iyilestirmeler": [iy("Dahiliye Polikliniği", "Poliklinik bekleme süresi", ["A-001", "A-004"], "Bekleme Süresi"),
                            iy("Ortopedi Polikliniği", "Bekleme ve bilgilendirme", ["A-028", "A-031"], "Bekleme Süresi"),
                            iy("Kadın Doğum Servisi", "Oda ve banyo temizliği", ["A-037", "A-043"], "Temizlik")]}
    if not dogru:
        t["iyilestirmeler"][0]["dayanak"].append("A-034")
        t["iyilestirmeler"].append(iy("Ortopedi Polikliniği", "Otopark", ["A-025"], ""))
        t["yorumlar"][10]["alinti"] = "Yemekler buz gibiydi"
    return t


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return taslak("<duzeltme>" in kullanici)


class KodTesti(unittest.TestCase):
    def test_okuma_ve_ozet(self):
        a, s = agent.anketleri_oku(ORNEK)
        self.assertEqual((len(a), s), (43, ["Hekim İlgisi", "Hemşire Bakımı", "Bekleme Süresi", "Temizlik"]))
        b = agent.ozet(a, s, "birim")
        self.assertEqual([(x["anahtar"], x["n"], x["genel"], x["memnun"]) for x in b],
                         [("Kadın Doğum Servisi", 10, 3.95, 80.0), ("Dahiliye Polikliniği", 24, 3.85, 67.7), ("Ortopedi Polikliniği", 9, 3.53, 58.3)])
        self.assertEqual(agent.dusukler(b, 4.0)["Dahiliye Polikliniği"], ["Hekim İlgisi", "Bekleme Süresi"])
        h = agent.ozet(a, s, "hekim")
        self.assertEqual([(x["anahtar"][1], x["n"]) for x in h][-1], ("Op. Dr. C. Yıldız", 9))
        self.assertEqual(agent.hekim_cikar("Kaya Bey çok ilgiliydi, Dr. kaya; Kayalar parkı", {"Uzm. Dr. A. Kaya"}), "[HEKİM] Bey çok ilgiliydi, Dr. [HEKİM]; Kayalar parkı")

    def test_denetim(self):
        a, s = agent.anketleri_oku(ORNEK)
        b = agent.ozet(a, s, "birim")
        d = agent.dusukler(b, 4.0)
        yorumlu = [x for x in a if x["yorum"]]
        hek = {x["hekim"] for x in a}
        self.assertEqual(agent.denetle(taslak(), yorumlu, d, hek), [])
        r = agent.denetle(taslak(False), yorumlu, d, hek)
        self.assertEqual(len(r), 3)
        self.assertTrue(any("A-034" in x and "başka birime" in x for x in r) and any("tek yoruma" in x for x in r) and any("A-040" in x for x in r))
        t = taslak(False)
        t["iyilestirmeler"][3]["dusuk_soru"] = "Bekleme Süresi"
        self.assertEqual(len(agent.denetle(t, yorumlu, d, hek)), 2)


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "m.xlsx"
            s = agent.calistir(ORNEK, cikti, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            for ad in ("Kaya", "Demir", "Yıldız", "Çelik"):
                self.assertNotIn(ad, MESAJLAR[0])
            self.assertIn("[HEKİM] Bey çok ilgiliydi", MESAJLAR[0])
            self.assertIn("düşük (< 4): Hemşire Bakımı, Bekleme Süresi, Temizlik", MESAJLAR[0])
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["İyileştirme Başlıkları", "Birim Özeti", "Hekim Özeti", "Yorumlar", "Tema Dağılımı"])
            h = wb["Hekim Özeti"]
            self.assertEqual((h.cell(5, 2).value, h.cell(5, 4).value), ("Op. Dr. C. Yıldız", None))
            self.assertIn("Yetersiz örnek", h.cell(5, 9).value)

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(ORNEK, Path(tmp) / "m.xlsx", tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 3)
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--girdi", str(ORNEK.with_name("yok.csv")), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
