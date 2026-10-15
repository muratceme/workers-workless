"""
Müşteri Yük Durum Bilgilendirmesi — Workers / Workless AI Agent
Karayolu Operasyon › Karayolu Operasyon Uzmanı

1. Kod sefer durum tablosunu okur, her sefer için gecikmeyi hesaplar (ETA − planlanan teslim; --tolerans dakikayı
   aşan gecikme, varsayılan 60) ve seferleri müşteriye göre gruplar. Plaka, şoför ve müşteri e-posta adresi modele
   gönderilmez.
2. Model her müşteri için tek bir yük durum e-postası yazar (konu + gövde): teslim edilen, zamanında ilerleyen ve
   geciken yükler; gecikmede yeni tahmini varış ve verideki gecikme nedeni.
3. Kod her e-postayı denetler:
   - Müşterinin bütün sefer ve yük numaraları geçiyor mu
   - Başka müşterinin sefer / yük numarası karışmış mı
   - ETA saatleri yazılmış mı
   - Veride olmayan bir sayı, saat ya da tarih eklenmiş mi
   - Gecikmesi olmayan müşteriye "gecikme" yazılmış mı
   - "Garanti" gibi taahhüt ifadesi var mı
   Sorun varsa yalnız o e-posta modele geri gönderilir (en fazla --tur kez).
4. E-postalar gönderilmez: Excel'de "Onay" sütunuyla ve müşteri başına .txt dosyası olarak hazırlanır.

Kullanım:
    python agent.py                                          # örnek sefer durumlarıyla dener
    python agent.py --girdi seferler.xlsx --firma "Örnek Lojistik" --iletisim "0212 000 00 00" --tolerans 30
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
SEMA = {
    "type": "object",
    "properties": {"epostalar": {"type": "array", "items": {"type": "object", "properties": {"musteri_no": S, "konu": S, "govde": S},
                                                            "required": ["musteri_no", "konu", "govde"], "additionalProperties": False}}},
    "required": ["epostalar"],
    "additionalProperties": False,
}
GECIKME = re.compile(r"\b(gecik|rotar|aksama|ertelen)")
TAAHHUT = re.compile(r"\b(garanti|kesinlikle teslim|soz veriyoruz|taahhut ediyoruz)")


def zaman_yaz(z: datetime | None) -> str:
    return z.strftime("%d.%m.%Y %H:%M") if z else ""


def seferleri_oku(yol: Path) -> list[dict]:
    r = kayitlar(yol, {"sefer": ("Sefer No", "Sefer"), "musteri": ("Müşteri", "Müşteri Adı"), "eposta": ("Müşteri E-posta", "E-posta"),
                       "yuk": ("Yük No", "İrsaliye No", "Sipariş No"), "cikis": ("Çıkış", "Yükleme Yeri"), "varis": ("Varış", "Teslim Yeri"),
                       "plan": ("Planlanan Teslim", "Teslim Tarihi"), "durum": ("Durum",), "konum": ("Son Konum", "Konum"),
                       "konum_zaman": ("Son Konum Zamanı", "Konum Zamanı"), "eta": ("ETA", "Tahmini Varış"), "neden": ("Gecikme Nedeni", "Açıklama")},
                 ("sefer", "musteri", "durum"))
    out = []
    for x in r:
        if not x["sefer"]:
            continue
        out.append({**{k: str(v).strip() for k, v in x.items() if k not in ("plan", "konum_zaman", "eta")},
                    "plan": zaman_coz(x["plan"]), "konum_zaman": zaman_coz(x["konum_zaman"]), "eta": zaman_coz(x["eta"])})
    return out


def degerlendir(s: dict, tolerans: int) -> tuple[str, int | None]:
    if katla(s["durum"]).startswith("teslim"):
        return "Teslim edildi", None
    if not s["eta"]:
        return "ETA bekleniyor", None
    if not s["plan"]:
        return "Planlanan teslim yok", None
    dk = int((s["eta"] - s["plan"]).total_seconds() // 60)
    return ("Gecikmeli" if dk > tolerans else "Zamanında"), dk


def musteri_gruplari(sefer: list[dict], tolerans: int) -> list[dict]:
    gr = defaultdict(list)
    for s in sefer:
        s["sonuc"], s["gecikme_dk"] = degerlendir(s, tolerans)
        gr[s["musteri"]].append(s)
    return [{"no": f"M{i}", "musteri": m, "eposta": next((s["eposta"] for s in l if s["eposta"]), ""), "seferler": l} for i, (m, l) in enumerate(gr.items(), 1)]


def olgular(g: dict) -> str:
    satir = []
    for s in g["seferler"]:
        satir.append(" | ".join([f"sefer {s['sefer']}", f"yük {s['yuk'] or '-'}", f"{s['cikis']} → {s['varis']}", f"planlanan teslim {zaman_yaz(s['plan']) or '-'}",
                                 f"durum {s['durum']}", f"son konum {s['konum'] or '-'} ({zaman_yaz(s['konum_zaman']) or '-'})",
                                 f"ETA {zaman_yaz(s['eta']) or 'yok'}", f"değerlendirme {s['sonuc']}" + (f" ({s['gecikme_dk']} dk)" if s["sonuc"] == "Gecikmeli" else ""),
                                 f"gecikme nedeni {s['neden'] or '-'}"]))
    return "\n".join(satir)


def denetle(gruplar: list[dict], sonuc: dict, firma: str, iletisim: str) -> list[str]:
    sorun = []
    tum = {g["no"]: {x for s in g["seferler"] for x in (s["sefer"], s["yuk"]) if x} for g in gruplar}
    for g in gruplar:
        e = sonuc.get(g["no"])
        if not e:
            sorun.append(f"{g['no']}: e-posta yazılmadı")
            continue
        metin = f"{e['konu']}\n{e['govde']}"
        for x in sorted(tum[g["no"]]):
            if x not in metin:
                sorun.append(f"{g['no']}: {x} e-postada geçmiyor")
        for no, l in tum.items():
            if no != g["no"]:
                for x in sorted(l):
                    if x in metin:
                        sorun.append(f"{g['no']}: başka müşterinin numarası ({x}) e-postada geçiyor")
        for s in g["seferler"]:
            if s["sonuc"] in ("Gecikmeli", "Zamanında") and s["eta"] and s["eta"].strftime("%H:%M") not in metin:
                sorun.append(f"{g['no']}: {s['sefer']} için tahmini varış saati {s['eta']:%H:%M} yazılmamış")
        izinli = sayilar(olgular(g) + " " + firma + " " + iletisim + " " + g["musteri"])
        fazla = sorted(sayilar(metin) - izinli)
        if fazla:
            sorun.append(f"{g['no']}: veride olmayan sayı / saat / tarih: {', '.join(fazla)}")
        if not any(s["sonuc"] == "Gecikmeli" for s in g["seferler"]) and GECIKME.search(katla(metin)):
            sorun.append(f"{g['no']}: bu müşterinin geciken yükü yok ama e-postada gecikme ifadesi var")
        if TAAHHUT.search(katla(metin)):
            sorun.append(f"{g['no']}: taahhüt ifadesi ('garanti' vb.) kullanılmış; tahmini olduğunu belirtin")
    return sorun


def uret(gruplar: list[dict], firma: str, iletisim: str, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    sonuc, hedef = {}, gruplar
    if duzeltme is not None:  # düzeltmede yalnız sorunlu e-postalar yeniden yazılır
        sorunlu = {s.split(":")[0] for s in duzeltme}
        sonuc.update({a: b for a, b in (onceki or {}).items() if a not in sorunlu})
        hedef = [g for g in gruplar if g["no"] in sorunlu]
    if not hedef:
        return sonuc
    bloklar = "\n".join(f'<musteri no="{g["no"]}" ad="{g["musteri"]}">\n{olgular(g)}\n</musteri>' for g in hedef)
    mesaj = f"<firma>{firma}</firma>\n<iletisim>{iletisim}</iletisim>\n<musteriler>\n{bloklar}\n</musteriler>"
    ids = {g["no"] for g in hedef}
    mesaj += duzeltme_blogu([s for s in (duzeltme or []) if s.split(":")[0] in ids], [v for a, v in (onceki or {}).items() if a in ids])
    for e in llm.json_iste(sistem, mesaj, SEMA).get("epostalar", []):
        if e.get("musteri_no") in ids:
            sonuc[e["musteri_no"]] = e
    return sonuc


def rapor_yaz(cikti: Path, gruplar, sonuc, kalan, tur) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "E-postalar"
    baslik(ws, ["Müşteri", "Alıcı", "Konu", "E-posta", "Seferler", "Geciken", "Onay"], (24, 26, 40, 90, 22, 10, 8))
    klasor = cikti.with_name(cikti.stem + "_epostalar")
    klasor.mkdir(parents=True, exist_ok=True)
    for g in gruplar:
        e = sonuc.get(g["no"])
        gec = sum(1 for s in g["seferler"] if s["sonuc"] == "Gecikmeli")
        ws.append([g["musteri"], g["eposta"], e["konu"] if e else "", e["govde"] if e else "AI yanıt vermedi", ", ".join(s["sefer"] for s in g["seferler"]), gec, ""])
        boya(ws, [3, 4], AI_DOLGU)
        if gec:
            boya(ws, [6], KIRMIZI)
        boya(ws, [7], SARI)
        if e:
            ad = re.sub(r"[^\w\-]+", "_", katla(g["musteri"]))[:40]
            (klasor / f"{g['no']}_{ad}.txt").write_text(f"Kime: {g['eposta']}\nKonu: {e['konu']}\n\n{e['govde']}\n", encoding="utf-8")
    son_not(ws, "Önemli: E-postalar gönderilmedi. Göndermeden önce sefer durumunu ve ETA'yı teyit edin.", kalan, tur)
    hizala(ws)
    d = wb.create_sheet("Sefer Durumu")
    baslik(d, ["Sefer No", "Müşteri", "Yük No", "Güzergâh", "Planlanan Teslim", "Durum", "Son Konum", "Konum Zamanı", "ETA", "Değerlendirme", "Gecikme (dk)",
               "Gecikme Nedeni"], (10, 24, 10, 20, 16, 12, 22, 16, 16, 14, 10, 40))
    for g in gruplar:
        for s in g["seferler"]:
            d.append([s["sefer"], s["musteri"], s["yuk"], f"{s['cikis']} → {s['varis']}", s["plan"], s["durum"], s["konum"], s["konum_zaman"], s["eta"],
                      s["sonuc"], s["gecikme_dk"], s["neden"]])
            for c in (5, 8, 9):
                d.cell(d.max_row, c).number_format = "DD.MM.YYYY HH:MM"
            boya(d, [10], KIRMIZI if s["sonuc"] == "Gecikmeli" else (SARI if s["sonuc"] in ("ETA bekleniyor", "Planlanan teslim yok") else YESIL))
    hizala(d)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    return klasor


def calistir(girdi: Path, cikti: Path, firma: str, iletisim: str, tolerans: int = 60, tur: int = 2, evet: bool = False) -> dict:
    sefer = seferleri_oku(girdi)
    if not sefer:
        raise llm.LLMHatasi(f"{girdi.name}: sefer yok.")
    gruplar = musteri_gruplari(sefer, tolerans)
    print(f"[OK] {len(sefer)} sefer · {len(gruplar)} müşteri · gecikmeli {sum(1 for s in sefer if s['sonuc'] == 'Gecikmeli')}")
    llm.onay_al("Müşteri adları ve sefer durumları (plaka, şoför ve e-posta adresi hariç) gönderilecek.", evet)
    sonuc, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(gruplar, firma, iletisim, d, o), lambda x: denetle(gruplar, x, firma, iletisim), tur)
    klasor = rapor_yaz(cikti, gruplar, sonuc, kalan, yapilan)
    return {"seferler": sefer, "gruplar": gruplar, "sonuc": sonuc, "kalan": kalan, "tur": yapilan, "klasor": klasor}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    p = argparse.ArgumentParser(description="Sefer durumundan müşterilere gönderilecek yük durum e-postalarını hazırlar (göndermez).")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "seferler.csv",
                   help="Sefer durumu (.xlsx/.csv): Sefer No, Müşteri, Durum; Yük No, Çıkış, Varış, Planlanan Teslim, Son Konum, ETA, Gecikme Nedeni …")
    p.add_argument("--firma", default="Örnek Lojistik", help="E-postayı imzalayacak firma adı")
    p.add_argument("--iletisim", default="0212 000 00 00", help="Müşterinin arayabileceği operasyon telefonu")
    p.add_argument("--tolerans", type=int, default=60, help="Bu kadar dakikayı aşan gecikme 'Gecikmeli' sayılır (varsayılan 60)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "yuk_durum.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.cikti, a.firma, a.iletisim, a.tolerans, a.tur, a.evet)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(s['sonuc'])}/{len(s['gruplar'])} e-posta hazır (gönderilmedi)")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()} · e-postalar: {s['klasor'].resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
