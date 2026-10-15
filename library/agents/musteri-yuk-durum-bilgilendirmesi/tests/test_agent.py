"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri" / "seferler.csv"
FIRMA, TEL = "Örnek Lojistik", "0212 000 00 00"
MESAJLAR = []
SON = f"\nSorularınız için {TEL} numaralı operasyon hattımızı arayabilirsiniz.\nSaygılarımızla,\n{FIRMA}"
DOGRU = {
    "M1": ("Yük durum bilgilendirmesi", "Sayın Örnek Gıda A.Ş. yetkilisi,\nSF-24520 (IRS-88122) 14.10.2026 17:20'de İzmir depoya teslim edilmiştir.\n"
           "SF-24518 (IRS-88120) Eskişehir çevre yolunda ilerliyor; tahmini varış 15.10.2026 13:30.\n"
           "SF-24519 (IRS-88121) Afyon girişinde kaza nedeniyle yol kapalı olduğu için gecikmektedir; yeni tahmini varış 15.10.2026 19:15." + SON),
    "M2": ("Yük durum bilgilendirmesi", "Sayın Mavi Yapı Malzemeleri Ltd. yetkilisi,\nSF-24530 (IRS-55410) Gebze depoda yüklendi; tahmini varış 16.10.2026 09:00.\n"
           "SF-24531 (IRS-55411) araç arızası nedeniyle gecikmektedir, yedek çekici yönlendirildi; yeni tahmini varış 15.10.2026 20:00." + SON),
    "M3": ("Yük durum bilgilendirmesi", "Sayın Delta Tekstil yetkilisi,\nSF-24540 (IRS-71001) Antalya-Burdur yolunda ilerliyor (son konum 15.10.2026 09:20). "
           "Tahmini varış saati netleştiğinde ayrıca bilgi vereceğiz." + SON),
}


def epostalar(dogru=True) -> dict:
    d = {k: {"musteri_no": k, "konu": a, "govde": b} for k, (a, b) in DOGRU.items()}
    if not dogru:
        d["M3"]["govde"] = d["M3"]["govde"].replace("ilerliyor", "ilerliyor, küçük bir gecikme olabilir")
        d["M1"]["govde"] = d["M1"]["govde"].replace(" (IRS-88122)", "")
        d["M2"]["govde"] += "\nNot: SF-24518 seferi de yoldadır."
    return d


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    d = epostalar("<duzeltme>" in kullanici)
    return {"epostalar": [d[n] for n in re.findall(r'<musteri no="([^"]+)"', kullanici)]}


class KodTesti(unittest.TestCase):
    def test_degerlendirme(self):
        g = agent.musteri_gruplari(agent.seferleri_oku(ORNEK), 60)
        self.assertEqual([(x["no"], x["musteri"], len(x["seferler"])) for x in g], [("M1", "Örnek Gıda A.Ş.", 3), ("M2", "Mavi Yapı Malzemeleri Ltd.", 2), ("M3", "Delta Tekstil", 1)])
        self.assertEqual([(s["sefer"], s["sonuc"], s["gecikme_dk"]) for x in g for s in x["seferler"]],
                         [("SF-24518", "Zamanında", -30), ("SF-24519", "Gecikmeli", 195), ("SF-24520", "Teslim edildi", None), ("SF-24530", "Zamanında", -60),
                          ("SF-24531", "Gecikmeli", 480), ("SF-24540", "ETA bekleniyor", None)])
        self.assertEqual(agent.musteri_gruplari(agent.seferleri_oku(ORNEK), 200)[0]["seferler"][1]["sonuc"], "Zamanında")
        o = agent.olgular(g[0])
        self.assertNotIn("16 ABC", o)
        self.assertNotIn("Ahmet", o)

    def test_denetim(self):
        g = agent.musteri_gruplari(agent.seferleri_oku(ORNEK), 60)
        self.assertEqual(agent.denetle(g, epostalar(), FIRMA, TEL), [])
        s = agent.denetle(g, epostalar(False), FIRMA, TEL)
        self.assertEqual(len(s), 4)  # sızan sefer no ayrıca "veride olmayan sayı" olarak da yakalanır
        self.assertIn("M2: veride olmayan sayı / saat / tarih: 24518", s)
        self.assertTrue(any("M1: IRS-88122" in x for x in s) and any("başka müşterinin numarası (SF-24518)" in x for x in s) and any("M3: bu müşterinin" in x for x in s))
        d = epostalar()
        d["M1"]["govde"] = d["M1"]["govde"].replace("19:15", "19:45") + " Teslimat garanti edilir."
        s = agent.denetle(g, d, FIRMA, TEL)
        self.assertEqual(len(s), 3)
        self.assertTrue(any("19:15" in x for x in s) and any("19:45" in x for x in s) and any("taahhüt" in x for x in s))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "y.xlsx"
            s = agent.calistir(ORNEK, cikti, FIRMA, TEL, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            for gizli in ("16 ABC 123", "Ahmet", "ornekgida.example"):
                self.assertNotIn(gizli, MESAJLAR[0])
            self.assertIn("değerlendirme Gecikmeli (195 dk)", MESAJLAR[0])
            self.assertEqual(re.findall(r'<musteri no="([^"]+)"', MESAJLAR[1]), ["M1", "M2", "M3"])
            ws = load_workbook(cikti)["E-postalar"]
            self.assertEqual((ws.cell(2, 2).value, ws.cell(2, 6).value, ws.cell(4, 6).value), ("lojistik@ornekgida.example", 1, 0))
            dosyalar = sorted(p.name for p in s["klasor"].iterdir())
            self.assertEqual(len(dosyalar), 3)
            self.assertTrue((s["klasor"] / dosyalar[0]).read_text(encoding="utf-8").startswith("Kime: lojistik@ornekgida.example"))

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(ORNEK, Path(tmp) / "y.xlsx", FIRMA, TEL, tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 4)
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--girdi", str(ORNEK.with_name("yok.csv")), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
