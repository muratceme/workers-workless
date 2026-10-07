"""
Mülakat Soru Seti Hazırlama — Workers / Workless AI Agent
İnsan Kaynakları › İşe Alım Uzmanı

1. Pozisyon bilgisi (talep formu / ilan) ve yetkinlik listesi (ad, tanım, ağırlık) okunur.
2. Model her yetkinlik için 2–3 yapılandırılmış soru yazar (davranışsal / durumsal / teknik): takip soruları,
   olumlu ve olumsuz göstergeler, davranışsal 1-3-5 puan tanımları ve tahmini süre; açılış ve kapanış metni.
3. Kod taslağı denetler: soru ve takip sorularında ayrımcı/kişisel hayata dair ifade (ortak tarayıcı), soru
   almamış yetkinlik, toplam sürenin mülakat süresini aşması. Sorun varsa taslak modele geri gönderilir
   (en fazla --tur kez).
4. Çıktı: Excel soru seti + formüllü değerlendirme formu (yetkinlik ortalaması, ağırlıklı toplam puan) + Markdown.

Kullanım:
    python agent.py                                              # örnek pozisyon ve yetkinliklerle dener
    python agent.py --pozisyon ilan.docx --yetkinlikler yetkinlikler.xlsx --sure 60
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import ayrimcilik
import belge
import llm

BURASI = Path(__file__).resolve().parent
ACILIS_KAPANIS_DK = 10


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def yetkinlikleri_oku(yol: Path | None) -> list[dict]:
    if not yol:
        return []
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    s = [r for r in s if any(c not in (None, "") for c in r)]
    b = [kucuk(x) for x in s[0]]
    i_ad = next((i for i, x in enumerate(b) if x in ("yetkinlik", "ad", "yetkinlik adı")), 0)
    i_tn = next((i for i, x in enumerate(b) if x in ("tanım", "açıklama", "davranış göstergesi")), None)
    i_ag = next((i for i, x in enumerate(b) if x in ("ağırlık", "ağırlık %", "önem")), None)
    liste = []
    for r in s[1:]:
        ad = str(r[i_ad] or "").strip()
        if ad:
            try:
                ag = float(str(r[i_ag]).replace(",", ".")) if i_ag is not None and r[i_ag] not in (None, "") else 1.0
            except ValueError:
                ag = 1.0
            liste.append({"ad": ad, "tanim": str(r[i_tn] or "").strip() if i_tn is not None else "", "agirlik": ag})
    return liste


def sema(yetkinlik_adlari: list[str]) -> dict:
    ad = {"type": "string", "enum": yetkinlik_adlari} if yetkinlik_adlari else {"type": "string"}
    return {
        "type": "object",
        "properties": {
            "acilis": {"type": "string"},
            "sorular": {"type": "array", "items": {
                "type": "object",
                "properties": {
                    "yetkinlik": ad,
                    "tur": {"type": "string", "enum": ["davranissal", "durumsal", "teknik"]},
                    "soru": {"type": "string"},
                    "takip_sorulari": {"type": "array", "items": {"type": "string"}},
                    "olumlu_gostergeler": {"type": "array", "items": {"type": "string"}},
                    "olumsuz_gostergeler": {"type": "array", "items": {"type": "string"}},
                    "puan_1": {"type": "string"}, "puan_3": {"type": "string"}, "puan_5": {"type": "string"},
                    "sure_dk": {"type": "integer"},
                },
                "required": ["yetkinlik", "tur", "soru", "takip_sorulari", "olumlu_gostergeler", "olumsuz_gostergeler",
                             "puan_1", "puan_3", "puan_5", "sure_dk"],
                "additionalProperties": False}},
            "kapanis": {"type": "string"},
            "yetkinlikler": {"type": "array", "items": {"type": "object", "properties": {"ad": {"type": "string"}, "tanim": {"type": "string"},
                                                                                         "agirlik": {"type": "number"}},
                                                        "required": ["ad", "tanim", "agirlik"], "additionalProperties": False}},
        },
        "required": ["acilis", "sorular", "kapanis", "yetkinlikler"],
        "additionalProperties": False,
    }


def denetle(t: dict, yetkinlikler: list[dict], sure: int) -> list[str]:
    sorunlar = []
    for i, s in enumerate(t["sorular"], 1):
        for b in ayrimcilik.tara(" ".join([s["soru"], *s["takip_sorulari"]])):
            if b["seviye"] == "yüksek" or b["kategori"] in ("Askerlik", "Kişisel veri"):
                sorunlar.append(f"Soru {i}: {b['kategori']} — '{b['ifade']}' ({b['aciklama']})")
    adlar = [y["ad"] for y in (yetkinlikler or t["yetkinlikler"])]
    for ad in adlar:
        if not any(s["yetkinlik"] == ad for s in t["sorular"]):
            sorunlar.append(f"'{ad}' yetkinliği için soru yok")
    toplam = sum(s["sure_dk"] for s in t["sorular"]) + ACILIS_KAPANIS_DK
    if toplam > sure:
        sorunlar.append(f"Toplam süre {toplam} dk (açılış-kapanış {ACILIS_KAPANIS_DK} dk dahil), mülakat süresi {sure} dk")
    return sorunlar


def uret(pozisyon: str, yetkinlikler: list[dict], sure: int, duzeltme: list[str] | None = None, onceki: dict | None = None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    y = ("\n".join(f"- {x['ad']} (ağırlık {x['agirlik']:g}): {x['tanim']}" for x in yetkinlikler)
         if yetkinlikler else "Verilmedi: pozisyona göre 4–6 temel yetkinlik belirle, tanımla ve ağırlıklandır (toplam 100); "
                              "'yetkinlikler' alanına yaz.")
    parca = [f"<pozisyon>\n{llm.maskele(pozisyon)}\n</pozisyon>", f"<yetkinlikler>\n{y}\n</yetkinlikler>", f"<sure>{sure}</sure>"]
    if duzeltme:
        parca.append("<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False) + "\nSorunlar:\n" + "\n".join(f"- {d}" for d in duzeltme)
                     + "\n</duzeltme>")
    t = llm.json_iste(sistem, "\n".join(parca), sema([x["ad"] for x in yetkinlikler]))
    if yetkinlikler:
        t["yetkinlikler"] = yetkinlikler            # kullanıcının listesi esastır
    return t


BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
GIRIS_DOLGU = PatternFill("solid", fgColor="FFF4CE")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
TUR_AD = {"davranissal": "Davranışsal", "durumsal": "Durumsal", "teknik": "Teknik"}


def rapor_yaz(cikti: Path, t: dict, sorunlar: list[str], sure: int, tur: int) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Soru Seti"
    ws.append(["No", "Yetkinlik", "Tür", "Soru", "Takip Soruları", "Olumlu Göstergeler", "Olumsuz Göstergeler", "1 — Zayıf", "3 — Yeterli",
               "5 — Güçlü", "Süre (dk)"])
    for h in ws[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for i, s in enumerate(t["sorular"], 1):
        ws.append([i, s["yetkinlik"], TUR_AD[s["tur"]], s["soru"], "\n".join(f"• {x}" for x in s["takip_sorulari"]),
                   "\n".join(f"+ {x}" for x in s["olumlu_gostergeler"]), "\n".join(f"− {x}" for x in s["olumsuz_gostergeler"]),
                   s["puan_1"], s["puan_3"], s["puan_5"], s["sure_dk"]])
        for h in ws[ws.max_row]:
            h.alignment = UST
    for j, w in enumerate((5, 22, 12, 50, 40, 36, 36, 28, 28, 28, 8), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "D2"

    # Değerlendirme formu: puan hücreleri sarı, yetkinlik ortalaması ve ağırlıklı toplam formülle
    f = wb.create_sheet("Değerlendirme Formu", 0)
    f.append(["Aday:", "", "Mülakatçı:", "", "Tarih:", ""])
    for c in (2, 4, 6):
        f.cell(1, c).fill = GIRIS_DOLGU
    f.append([])
    f.append(["No", "Yetkinlik", "Soru", "Puan (1–5)", "Notlar / STAR kanıtı"])
    for h in f[3]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    ilk = f.max_row + 1
    for i, s in enumerate(t["sorular"], 1):
        f.append([i, s["yetkinlik"], s["soru"], None, ""])
        f.cell(f.max_row, 4).fill = GIRIS_DOLGU
        f.cell(f.max_row, 5).fill = GIRIS_DOLGU
        for h in f[f.max_row]:
            h.alignment = UST
    son = f.max_row
    f.append([])
    f.append(["", "Yetkinlik", "Ağırlık", "Ortalama Puan", "Ağırlıklı Katkı"])
    for h in f[f.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    ozet_ilk = f.max_row + 1
    for y in t["yetkinlikler"]:
        r = f.max_row + 1
        f.append(["", y["ad"], y["agirlik"],
                  f'=IFERROR(AVERAGEIF($B${ilk}:$B${son},B{r},$D${ilk}:$D${son}),"")',
                  f'=IF(D{r}="","",C{r}*D{r})'])
    ozet_son = f.max_row
    f.append(["", "TOPLAM (5 üzerinden)", f"=SUM(C{ozet_ilk}:C{ozet_son})", "",
              # Yalnız puanlanmış yetkinliklerin ağırlıkları paydaya girer ("" döndüren formüller hariç)
              f'=IFERROR(SUM(E{ozet_ilk}:E{ozet_son})/SUMPRODUCT((D{ozet_ilk}:D{ozet_son}<>"")*C{ozet_ilk}:C{ozet_son}),"")'])
    for h in f[f.max_row]:
        h.font = Font(bold=True)
    f.cell(f.max_row, 5).number_format = "0.00"
    for r in range(ozet_ilk, ozet_son + 1):
        f.cell(r, 4).number_format = "0.00"
        f.cell(r, 5).number_format = "0.00"
    f.append([])
    f.append(["", "Genel değerlendirme / karar önerisi:"])
    f.cell(f.max_row, 3).fill = GIRIS_DOLGU
    for j, w in enumerate((5, 30, 70, 14, 50), 1):
        f.column_dimensions[get_column_letter(j)].width = w

    a = wb.create_sheet("Açılış ve Kapanış")
    for satir in [["Açılış", t["acilis"]], ["Kapanış", t["kapanis"]], [],
                  ["Sorulmaması gerekenler", "Yaş, medeni hâl, çocuk/aile planı, gebelik, sağlık ve engellilik, din ve inanç, etnik köken ve "
                                             "memleket, siyasi görüş, sendika üyeliği, askerlik, görünüş ve kilo (4857 s. Kanun md. 5; 6701 s. "
                                             "Kanun md. 3 ve 6; 6356 s. Kanun md. 25)"],
                  ["Son kontrol", "; ".join(sorunlar) or "Ayrımcı soru yok, tüm yetkinlikler kapsandı, süre uygun"],
                  ["Süre", f"{sum(s['sure_dk'] for s in t['sorular']) + ACILIS_KAPANIS_DK} / {sure} dk (açılış-kapanış {ACILIS_KAPANIS_DK} dk dahil)"],
                  ["Düzeltme turu", tur], ["Model", llm.kullanim_ozeti()], ["Oluşturulma", datetime.now().strftime("%d.%m.%Y %H:%M")]]:
        a.append(satir)
    a.cell(5, 2).fill = KIRMIZI if sorunlar else YESIL
    a.column_dimensions["A"].width = 22
    a.column_dimensions["B"].width = 120
    for satir in a.iter_rows():
        for h in satir:
            h.alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

    md = ["# Mülakat soru seti", "", "## Açılış", t["acilis"], ""]
    for y in t["yetkinlikler"]:
        md += [f"## {y['ad']} (ağırlık {y['agirlik']:g})", y["tanim"], ""]
        for s in [s for s in t["sorular"] if s["yetkinlik"] == y["ad"]]:
            md += [f"**{TUR_AD[s['tur']]}** — {s['soru']} _({s['sure_dk']} dk)_", *[f"- {x}" for x in s["takip_sorulari"]],
                   f"- 1: {s['puan_1']} · 3: {s['puan_3']} · 5: {s['puan_5']}", ""]
    md += ["## Kapanış", t["kapanis"]]
    yol = cikti.with_suffix(".md")
    yol.write_text("\n".join(md), encoding="utf-8")
    return yol


def calistir(pozisyon_yolu: Path, cikti: Path, yetkinlik_yolu: Path | None = None, sure: int = 60, tur: int = 2, evet: bool = False) -> dict:
    pozisyon = belge.metin_oku(pozisyon_yolu)
    yetkinlikler = yetkinlikleri_oku(yetkinlik_yolu)
    print(f"[OK] Pozisyon bilgisi okundu · {len(yetkinlikler) or 'model belirleyecek'} yetkinlik · {sure} dk")
    llm.onay_al("Pozisyon bilgisi (e-posta/telefon maskeli) ve yetkinlik listesi soru seti için gönderilecek.", evet)
    t = uret(pozisyon, yetkinlikler, sure)
    sorunlar = denetle(t, yetkinlikler, sure)
    yapilan = 0
    while sorunlar and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(sorunlar)} sorun modele geri gönderiliyor")
        t = uret(pozisyon, yetkinlikler, sure, sorunlar, t)
        sorunlar = denetle(t, yetkinlikler, sure)
    md = rapor_yaz(cikti, t, sorunlar, sure, yapilan)
    return {"taslak": t, "sorunlar": sorunlar, "tur": yapilan, "md": md}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Pozisyon ve yetkinliklerden yapılandırılmış mülakat soru seti hazırlar.")
    p.add_argument("--pozisyon", type=Path, default=BURASI / "ornek_veri" / "pozisyon_talep_formu.txt", help="Talep formu veya ilan (.txt/.md/.docx/.pdf)")
    p.add_argument("--yetkinlikler", type=Path, help="Yetkinlik listesi (.xlsx/.csv: Yetkinlik, Tanım, Ağırlık); verilmezse model önerir")
    p.add_argument("--sure", type=int, default=60, help="Mülakat süresi, dakika (varsayılan 60)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "mulakat_soru_seti.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    if a.pozisyon == BURASI / "ornek_veri" / "pozisyon_talep_formu.txt":
        a.yetkinlikler = a.yetkinlikler or BURASI / "ornek_veri" / "yetkinlikler.csv"
    try:
        s = calistir(a.pozisyon, a.cikti, a.yetkinlikler, a.sure, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(s['taslak']['sorular'])} soru · düzeltme turu {s['tur']}")
    for x in s["sorunlar"]:
        print(f"[!] {x}")
    print(f"[OK] Rapor: {a.cikti.resolve()} (+ {s['md'].name})")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
