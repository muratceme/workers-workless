"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
YETKINLIK = agent.yetkinlikleri_oku(ORNEK / "yetkinlikler.csv")
MESAJLAR = []


def soru(y, metin, dk=6):
    return {"yetkinlik": y, "tur": "davranissal", "soru": metin, "takip_sorulari": ["Sonuç ne oldu?"], "olumlu_gostergeler": ["somut örnek"],
            "olumsuz_gostergeler": ["genel ifade"], "puan_1": "örnek yok", "puan_3": "örnek var", "puan_5": "ölçülebilir sonuç", "sure_dk": dk}


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    adlar = [y["ad"] for y in YETKINLIK]
    if "<duzeltme>" not in kullanici:     # ilk taslak hatalı: medeni hâl sorusu, eksik yetkinlik, süre aşımı
        sorular = [soru(adlar[0], "Bir itirazı nasıl yönettiğinizi anlatır mısınız?", 20),
                   soru(adlar[1], "Evli misiniz, sık seyahat sorun olur mu?", 20),
                   soru(adlar[2], "Haftalık rotanızı nasıl planlarsınız?", 20)]
    else:
        sorular = [soru(a, f"{a} ile ilgili bir deneyiminizi anlatır mısınız?") for a in adlar]
    return {"acilis": "Hoş geldiniz.", "sorular": sorular, "kapanis": "Teşekkürler.", "yetkinlikler": []}


class OkumaTesti(unittest.TestCase):
    def test_yetkinlikler(self):
        self.assertEqual([y["agirlik"] for y in YETKINLIK], [30, 25, 20, 15, 10])
        self.assertEqual(YETKINLIK[0]["ad"], "Müşteri ilişkileri ve ikna")

    def test_denetim(self):
        t = sahte_json_iste("", "", {})
        s = agent.denetle(t, YETKINLIK, 60)
        self.assertTrue(any("Medeni hâl" in x for x in s))
        self.assertTrue(any("Analitik düşünme" in x for x in s))
        self.assertTrue(any("70 dk" in x for x in s))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_ve_form(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "m.xlsx"
            s = agent.calistir(ORNEK / "pozisyon_talep_formu.txt", cikti, ORNEK / "yetkinlikler.csv", 60, evet=True)
            self.assertEqual((len(MESAJLAR), s["tur"], s["sorunlar"]), (2, 1, []))
            self.assertIn("Evli misiniz", MESAJLAR[1])                      # düzeltme isteğinde önceki taslak gider
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Değerlendirme Formu", "Soru Seti", "Açılış ve Kapanış"])
            f = wb["Değerlendirme Formu"]
            formuller = [c.value for r in f.iter_rows() for c in r if isinstance(c.value, str) and c.value.startswith("=")]
            self.assertTrue(any("AVERAGEIF" in x for x in formuller))
            self.assertTrue(any("SUMPRODUCT" in x for x in formuller))
            self.assertEqual(len([x for x in formuller if "AVERAGEIF" in x]), 5)
            self.assertIn("Analitik düşünme", s["md"].read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
