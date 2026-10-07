"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import re
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))

import os  # noqa: E402
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz
import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    bolum = kullanici.split("<baglam_icin_diger_acik_kalemler>")[0]
    etiketler = re.findall(r"^([BD]\d{4}) \|", bolum, re.M)
    return {
        "kalemler": [{"etiket": e, "olasi_neden": "Diğer", "iliskili_kalem": "", "aciklama": "test",
                      "onerilen_kayit": "Kayıt gerekmez", "guven": "orta"} for e in etiketler]
        + [{"etiket": "B9999", "olasi_neden": "Diğer", "iliskili_kalem": "", "aciklama": "uydurma",
            "onerilen_kayit": "", "guven": "dusuk"}],
        "genel_degerlendirme": "Genel test yorumu.",
    }


class AgentTesti(unittest.TestCase):
    def setUp(self):
        MESAJLAR.clear()

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_yalnizca_acik_kalemler_gonderilir(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "m.xlsx"
            e, ab, ad, sonuc = agent.calistir(ORNEK / "banka_ekstresi.csv", ORNEK / "defter_102.csv", cikti, evet=True)

            self.assertEqual((len(e), len(ab), len(ad)), (9, 6, 4))
            self.assertEqual(len(MESAJLAR), 1)
            self.assertNotIn("Kira ödemesi", MESAJLAR[0])      # eşleşen kayıt gönderilmez
            self.assertIn("EFT komisyonu", MESAJLAR[0])
            self.assertEqual(len(sonuc), 10)
            self.assertNotIn("B9999", sonuc)                     # uydurma numara yok sayılır

            wb = load_workbook(cikti)
            self.assertIn("AI Değerlendirme", wb.sheetnames)
            ws = wb["Açık - Banka"]
            basliklar = [h.value for h in ws[1]]
            self.assertEqual(basliklar[-1], "İnsan Onayı")
            self.assertEqual(ws.cell(2, basliklar.index("AI Nedeni") + 1).value, "Diğer")

    def test_iban_maskelenir(self):
        k = agent.cekirdek.Kayit("banka", 1, agent.cekirdek.tarih_coz("2026-09-01"),
                                 "EFT TR33 0006 1005 1978 6457 8413 26", -10.0)
        self.assertIn("[IBAN]", agent.kalem_satiri(k, [], [], 3))


if __name__ == "__main__":
    unittest.main()
