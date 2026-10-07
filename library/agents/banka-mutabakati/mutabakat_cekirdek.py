"""
Banka Mutabakatı — Workers / Workless kod bloğu
Finans & Muhasebe › Muhasebe Elemanı

Banka ekstresi ile muhasebe defterini (102 Bankalar muavini) eşleştirir; eşleşen
kayıtları, iki taraftaki açık kalemleri ve kural tabanlı olası nedenleri tek bir
Excel raporunda verir. İnternete bağlanmaz, veriniz bilgisayarınızdan çıkmaz.

Eşleştirme sırası:
  1. Referans  : açıklamalarda ortak dekont/fatura numarası (6+ hane) ve aynı tutar
  2. Tutar+Tarih: aynı tutar ve tarih farkı ≤ tolerans (en yakın tarih, sonra en benzer açıklama)

Kullanım:
    python main.py                                         # örnek veriyle dener
    python main.py --banka ekstre.xlsx --defter muavin.xlsx --cikti cikti/mutabakat.xlsx
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from dataclasses import dataclass
from datetime import date, datetime
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent

# ----------------------------------------------------------------------------
# Okuma ve normalleştirme
# ----------------------------------------------------------------------------

_TR_ASCII = str.maketrans("çğıöşüÇĞİÖŞÜI", "cgiosuCGIOSUI")

SUTUN_ADLARI = {
    "tarih": {"tarih", "islem tarihi", "valor", "date", "fis tarihi", "belge tarihi"},
    "aciklama": {"aciklama", "islem aciklamasi", "description", "detay", "fis aciklamasi"},
    "tutar": {"tutar", "islem tutari", "amount", "net tutar"},
    "borc": {"borc", "borc tutari", "debit"},
    "alacak": {"alacak", "alacak tutari", "credit"},
}


def _anahtar(s: str) -> str:
    return re.sub(r"\s+", " ", str(s or "").translate(_TR_ASCII).lower().strip())


@dataclass
class Kayit:
    kaynak: str          # "banka" | "defter"
    sira: int            # dosyadaki satır no (başlık hariç, 1'den)
    tarih: date
    aciklama: str
    tutar: float         # giriş +, çıkış −

    @property
    def etiket(self) -> str:
        return f"{self.kaynak[0].upper()}{self.sira:04d}"


def sayi_coz(deger) -> float:
    """'1.234,56' · '1,234.56' · '-1234.5' · 1234.5 → float."""
    if deger is None or deger == "":
        return 0.0
    if isinstance(deger, (int, float)):
        return float(deger)
    s = str(deger).strip().replace(" ", "").replace("TL", "").replace("₺", "")
    negatif = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    deger = float(s)
    return -deger if negatif else deger


def tarih_coz(deger) -> date:
    if isinstance(deger, datetime):
        return deger.date()
    if isinstance(deger, date):
        return deger
    s = str(deger).strip()
    for bicim in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s, bicim).date()
        except ValueError:
            continue
    raise ValueError(f"Tarih anlaşılamadı: {s!r}")


def _satirlari_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        ws = load_workbook(yol, data_only=True, read_only=True).active
        return [list(r) for r in ws.iter_rows(values_only=True)]
    metin = None
    for kodlama in ("utf-8-sig", "cp1254"):
        try:
            metin = yol.read_text(encoding=kodlama)
            break
        except UnicodeDecodeError:
            continue
    ayirici = csv.Sniffer().sniff(metin.splitlines()[0], delimiters=",;\t").delimiter
    return list(csv.reader(metin.splitlines(), delimiter=ayirici))


def dosya_oku(yol: Path, kaynak: str) -> list[Kayit]:
    """
    Standart sütunlar: tarih, aciklama ve ya `tutar` (giriş +, çıkış −)
    ya da `borc` + `alacak`.
      banka : tutar = alacak − borç   (ekstrede alacak = hesaba giriş)
      defter: tutar = borç − alacak   (102 hesapta borç = bankaya giriş)
    """
    satirlar = [r for r in _satirlari_oku(yol) if any(c not in (None, "") for c in r)]
    if not satirlar:
        raise SystemExit(f"{yol} boş.")
    basliklar = [_anahtar(b) for b in satirlar[0]]
    konum = {}
    for alan, adlar in SUTUN_ADLARI.items():
        for i, b in enumerate(basliklar):
            if b in adlar:
                konum[alan] = i
                break
    eksik = [a for a in ("tarih", "aciklama") if a not in konum]
    if eksik or not ("tutar" in konum or {"borc", "alacak"} <= konum.keys()):
        raise SystemExit(
            f"{yol.name}: sütunlar tanınmadı. Gerekli: tarih, aciklama ve tutar (ya da borc + alacak). "
            f"Bulunan başlıklar: {satirlar[0]}"
        )

    kayitlar = []
    for n, r in enumerate(satirlar[1:], start=1):
        if "tutar" in konum:
            tutar = sayi_coz(r[konum["tutar"]])
        else:
            borc, alacak = sayi_coz(r[konum["borc"]]), sayi_coz(r[konum["alacak"]])
            tutar = (alacak - borc) if kaynak == "banka" else (borc - alacak)
        kayitlar.append(Kayit(kaynak, n, tarih_coz(r[konum["tarih"]]), str(r[konum["aciklama"]] or "").strip(), round(tutar, 2)))
    return kayitlar


# ----------------------------------------------------------------------------
# Eşleştirme
# ----------------------------------------------------------------------------

REFERANS = re.compile(r"\b[A-Z]{0,4}\d{6,}\b")


def referanslar(aciklama: str) -> set[str]:
    return set(REFERANS.findall(aciklama.upper()))


def benzerlik(a: str, b: str) -> float:
    return SequenceMatcher(None, _anahtar(a), _anahtar(b)).ratio()


@dataclass
class Eslesme:
    banka: Kayit
    defter: Kayit
    yontem: str

    @property
    def gun_farki(self) -> int:
        return (self.defter.tarih - self.banka.tarih).days


def eslestir(banka: list[Kayit], defter: list[Kayit], gun_toleransi: int = 3, tutar_toleransi: float = 0.01):
    """(eşleşmeler, açık banka kalemleri, açık defter kalemleri) döndürür."""
    acik_defter = list(defter)
    eslesmeler: list[Eslesme] = []
    kalan_banka: list[Kayit] = []

    # 1) referans
    for b in banka:
        refs = referanslar(b.aciklama)
        aday = next((d for d in acik_defter if refs and refs & referanslar(d.aciklama)
                     and abs(d.tutar - b.tutar) <= tutar_toleransi), None)
        if aday:
            eslesmeler.append(Eslesme(b, aday, "Referans"))
            acik_defter.remove(aday)
        else:
            kalan_banka.append(b)

    # 2) tutar + tarih
    acik_banka = []
    for b in sorted(kalan_banka, key=lambda k: k.tarih):
        adaylar = [d for d in acik_defter
                   if abs(d.tutar - b.tutar) <= tutar_toleransi and abs((d.tarih - b.tarih).days) <= gun_toleransi]
        if adaylar:
            en_iyi = min(adaylar, key=lambda d: (abs((d.tarih - b.tarih).days), -benzerlik(b.aciklama, d.aciklama)))
            eslesmeler.append(Eslesme(b, en_iyi, "Tutar+Tarih"))
            acik_defter.remove(en_iyi)
        else:
            acik_banka.append(b)
    return eslesmeler, acik_banka, acik_defter


# ----------------------------------------------------------------------------
# Kural tabanlı olası neden
# ----------------------------------------------------------------------------

MASRAF = re.compile(r"komisyon|masraf|ucret|bsmv|kesinti|aidat|hesap isletim|eft ucret|havale ucret", re.I)
FAIZ = re.compile(r"faiz|nemalandirma|getiri|repo", re.I)


def tl(x: float) -> str:
    """1234.5 → '1.234,50'"""
    return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def olasi_neden(k: Kayit, karsi_taraf: list[Kayit], ayni_taraf: list[Kayit], gun_toleransi: int) -> str:
    acik = _anahtar(k.aciklama)
    if k.kaynak == "banka" and MASRAF.search(acik):
        return "Banka masrafı/komisyonu: deftere işlenmemiş olabilir"
    if k.kaynak == "banka" and FAIZ.search(acik):
        return "Faiz/getiri: deftere işlenmemiş olabilir"
    ayni_tutar = [x for x in karsi_taraf if abs(x.tutar - k.tutar) <= 0.01]
    if ayni_tutar:
        fark = min(abs((x.tarih - k.tarih).days) for x in ayni_tutar)
        return f"Karşı tarafta aynı tutar {fark} gün farkla var: zamanlama farkı olabilir (tolerans {gun_toleransi} gün)"
    yakin = [x for x in karsi_taraf if x.tutar and abs(x.tutar - k.tutar) / abs(x.tutar) < 0.05]
    if yakin:
        return f"Karşı tarafta benzer tutar var ({tl(yakin[0].tutar)}): tutar hatası/hane kayması olabilir"
    ikiz = [x for x in ayni_taraf if x is not k and abs(x.tutar - k.tutar) <= 0.01 and benzerlik(x.aciklama, k.aciklama) > 0.8]
    if ikiz:
        return f"Aynı tutar ve açıklamayla {ikiz[0].etiket} kaydı da var: mükerrer kayıt olabilir"
    if k.kaynak == "defter":
        return "Defterde olup bankada yok: bankaya henüz yansımamış işlem olabilir"
    return "Bankada olup defterde yok: kaydı atlanmış olabilir"


# ----------------------------------------------------------------------------
# Excel çıktısı
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
PARA = '#,##0.00;[Red]-#,##0.00'


def _sayfa(wb, ad, basliklar, satirlar, genislikler, para_sutunlari=()):
    ws = wb.create_sheet(ad)
    ws.append(basliklar)
    for h in ws[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for s in satirlar:
        ws.append(s)
    for i, g in enumerate(genislikler, 1):
        ws.column_dimensions[get_column_letter(i)].width = g
    for sutun in para_sutunlari:
        for hucre in ws[get_column_letter(sutun)][1:]:
            hucre.number_format = PARA
    ws.freeze_panes = "A2"
    if satirlar:
        ws.auto_filter.ref = ws.dimensions
    return ws


def acik_kalem_satirlari(kalemler, karsi_taraf, ayni_taraf, gun_toleransi):
    return [[k.etiket, k.tarih, k.aciklama, k.tutar, olasi_neden(k, karsi_taraf, ayni_taraf, gun_toleransi)] for k in kalemler]


def rapor_yaz(cikti: Path, banka, defter, eslesmeler, acik_banka, acik_defter, gun_toleransi: int) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    toplam_b, toplam_d = sum(k.tutar for k in banka), sum(k.tutar for k in defter)
    ozet = [
        ["Banka hareket sayısı", len(banka)],
        ["Defter kayıt sayısı", len(defter)],
        ["Eşleşen", len(eslesmeler)],
        ["  · Referans ile", sum(1 for e in eslesmeler if e.yontem == "Referans")],
        ["  · Tutar+Tarih ile", sum(1 for e in eslesmeler if e.yontem == "Tutar+Tarih")],
        ["Açık kalem: banka", len(acik_banka)],
        ["Açık kalem: defter", len(acik_defter)],
        ["Banka hareket toplamı", round(toplam_b, 2)],
        ["Defter hareket toplamı", round(toplam_d, 2)],
        ["Fark (banka − defter)", round(toplam_b - toplam_d, 2)],
        ["Açık kalemlerle açıklanan fark", round(sum(k.tutar for k in acik_banka) - sum(k.tutar for k in acik_defter), 2)],
    ]
    ws = _sayfa(wb, "Özet", ["Gösterge", "Değer"], ozet, [34, 18])
    for r in range(9, 13):
        ws.cell(r, 2).number_format = PARA

    _sayfa(wb, "Eşleşenler",
           ["Banka No", "Banka Tarih", "Banka Açıklama", "Tutar", "Defter No", "Defter Tarih", "Defter Açıklama", "Gün Farkı", "Yöntem"],
           [[e.banka.etiket, e.banka.tarih, e.banka.aciklama, e.banka.tutar, e.defter.etiket, e.defter.tarih,
             e.defter.aciklama, e.gun_farki, e.yontem] for e in eslesmeler],
           [10, 12, 40, 14, 10, 12, 40, 10, 12], para_sutunlari=(4,))
    sutunlar = ["No", "Tarih", "Açıklama", "Tutar", "Olası Neden (kural tabanlı)"]
    _sayfa(wb, "Açık - Banka", sutunlar, acik_kalem_satirlari(acik_banka, acik_defter, banka, gun_toleransi), [10, 12, 44, 14, 70], (4,))
    _sayfa(wb, "Açık - Defter", sutunlar, acik_kalem_satirlari(acik_defter, acik_banka, defter, gun_toleransi), [10, 12, 44, 14, 70], (4,))
    for ws in wb.worksheets[1:]:
        for satir in ws.iter_rows(min_row=2):
            for h in satir:
                if isinstance(h.value, date):
                    h.number_format = "DD.MM.YYYY"

    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


# ----------------------------------------------------------------------------
# Komut satırı
# ----------------------------------------------------------------------------

def calistir(banka_yolu: Path, defter_yolu: Path, cikti: Path, gun_toleransi: int = 3, tutar_toleransi: float = 0.01):
    banka = dosya_oku(banka_yolu, "banka")
    defter = dosya_oku(defter_yolu, "defter")
    eslesmeler, acik_banka, acik_defter = eslestir(banka, defter, gun_toleransi, tutar_toleransi)
    rapor_yaz(cikti, banka, defter, eslesmeler, acik_banka, acik_defter, gun_toleransi)
    return eslesmeler, acik_banka, acik_defter


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Banka ekstresi ile defter kayıtlarını eşleştirir.")
    p.add_argument("--banka", type=Path, default=BURASI / "ornek_veri" / "banka_ekstresi.csv", help="Banka ekstresi (.csv/.xlsx)")
    p.add_argument("--defter", type=Path, default=BURASI / "ornek_veri" / "defter_102.csv", help="Defter / 102 muavini (.csv/.xlsx)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "mutabakat.xlsx", help="Excel çıktı yolu")
    p.add_argument("--gun-toleransi", type=int, default=3, help="Tutar+Tarih eşleşmesinde izin verilen gün farkı")
    p.add_argument("--tutar-toleransi", type=float, default=0.01, help="Kuruş yuvarlama toleransı")
    a = p.parse_args(argv)

    eslesmeler, acik_b, acik_d = calistir(a.banka, a.defter, a.cikti, a.gun_toleransi, a.tutar_toleransi)
    print(f"[OK] {len(eslesmeler)} kayıt eşleşti · açık kalem: banka {len(acik_b)}, defter {len(acik_d)}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
