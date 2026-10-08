"""
Dilekçe Taslağı — Workers / Workless AI Agent
Hukuk Bürosu › Dava Takip › Avukat

1. Klasördeki dava bilgileri ('Alan: değer' satırları), olay özeti, talepler ve deliller listesi okunur. Varsa
   mevzuat/ klasöründeki kanun maddesi metinleri maddelere bölünür.
2. Kod HMK md. 119'da sayılan unsurları kontrol eder: mahkeme, taraf adları ve adresleri, davacı TCKN (gerçek kişi)
   veya vergi no (tüzel kişi, kontrol hanesiyle), vekil, konu, dava değeri, olay özeti, deliller, talepler. Dava
   türü zorunlu arabuluculuğa tabi olabilecek türdense arabuluculuk son tutanağının eklenip eklenmediğine bakar.
3. Taraf ve vekil adları [DAVACI] / [DAVALI] / [DAVACI VEKİLİ], ek adlar [KİŞİ-n] takma adlarıyla; telefon,
   e-posta, TCKN ve IBAN maskelenir. Adresler ve kimlik numaraları modele gönderilmez.
4. Model açıklamaları (her vakıa için olay özetinden birebir dayanak alıntısı ve delil numaralarıyla), hukuki
   sebepleri, sonuç ve istemi ve avukata notları yazar.
5. Kod modelin çıktısını denetler: dayanak alıntısı olay özetinde birebir geçiyor mu, delil numaraları var mı,
   delilsiz vakıa ve kullanılmayan delil, verilmeyen maddeye atıf veya listede olmayan kanun sayısı, girdilerde
   olmayan tarih / tutar, dava değeri sonuç ve istemde geçiyor mu.
6. Çıktı: Markdown dilekçe taslağı (başlık, taraflar, açıklamalar, hukuki sebepler, deliller, sonuç ve istem,
   ekler; takma adlar geri açılmış) + Excel (Kontrol Listesi, Vakıa–Delil, Hukuki Sebepler, Kontroller, Avukata
   Notlar).

Kullanım:
    python agent.py                                          # örnek: kurgusal ticari alacak davası
    python agent.py --girdi ./dava_klasoru --gizle "Tanık Adı"
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
KANUNLAR = {"4721": "Türk Medeni Kanunu", "6098": "Türk Borçlar Kanunu", "6102": "Türk Ticaret Kanunu", "6100": "Hukuk Muhakemeleri Kanunu",
            "2004": "İcra ve İflas Kanunu", "3095": "Kanuni Faiz ve Temerrüt Faizine İlişkin Kanun",
            "6325": "Hukuk Uyuşmazlıklarında Arabuluculuk Kanunu", "4857": "İş Kanunu", "7036": "İş Mahkemeleri Kanunu",
            "6502": "Tüketicinin Korunması Hakkında Kanun", "492": "Harçlar Kanunu"}
# Zorunlu arabuluculuğa tabi olabilecek uyuşmazlık türlerinin anahtar kelimeleri (avukat kontrolü için ipucu)
ARABULUCULUK_IPUCU = ("ticari", "isci", "is davasi", "kidem", "ihbar", "iscilik", "tuketici", "kira", "ortakligin giderilmesi", "kat mulkiyeti",
                      "komsu")
ALANLAR = {"mahkeme": ("mahkeme",), "davaci": ("davaci", "davaci adi", "davaci unvani"),
           "davaci_tckn": ("davaci tckn", "davaci tc kimlik no", "davaci t c kimlik no", "tckn"),
           "davaci_vkn": ("davaci vergi no", "davaci vkn", "davaci vergi kimlik no"), "davaci_mersis": ("davaci mersis no", "mersis no"),
           "davaci_adres": ("davaci adres", "davaci adresi"), "vekil": ("davaci vekili", "vekil", "vekili"), "vekil_adres": ("vekil adres", "vekil adresi"),
           "davali": ("davali", "davali adi", "davali unvani"), "davali_tckn": ("davali tckn",), "davali_vkn": ("davali vergi no", "davali vkn"),
           "davali_adres": ("davali adres", "davali adresi"), "tur": ("dava turu", "tur"), "deger": ("dava degeri", "deger"),
           "konu": ("konu", "dava konusu")}
TUZEL = ("a s", "as", "ltd", "sti", "anonim", "limited", "koop", "kooperatif", "vakfi", "dernegi", "belediyesi", "bakanligi", "sirketi")


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def katla(s) -> str:
    s = kucuk(s)
    for a, b in zip("ıöüşğçâîû", "iousgcaiu"):
        s = s.replace(a, b)
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def buyuk(s: str) -> str:
    return str(s).replace("i", "İ").replace("ı", "I").upper()


def tckn_gecerli(no: str) -> bool:
    if not re.fullmatch(r"[1-9]\d{10}", no):
        return False
    d = [int(c) for c in no]
    return d[9] == (sum(d[0:9:2]) * 7 - sum(d[1:8:2])) % 10 and d[10] == sum(d[:10]) % 10


def vkn_gecerli(no: str) -> bool:
    if not re.fullmatch(r"\d{10}", no):
        return False
    s = 0
    for i, n in enumerate(reversed(no[:9]), 1):
        c1 = (int(n) + i) % 10
        if c1:
            s += (c1 * 2 ** i) % 9 or 9
    return (10 - s % 10) % 10 == int(no[9])


def tuzel_mi(ad: str) -> bool:
    k = f" {katla(ad)} "
    return any(f" {t} " in k for t in TUZEL)


# ----------------------------------------------------------------------------
# Girdiler
# ----------------------------------------------------------------------------

@dataclass
class Delil:
    no: str
    ad: str
    tur: str
    ek: str
    vakialar: list = field(default_factory=list)


@dataclass
class Madde:
    id: str
    kaynak: str
    no: str
    metin: str


@dataclass
class Dava:
    bilgi: dict
    olay: str
    talepler: list[str]
    deliller: list[Delil]
    maddeler: list[Madde]
    kaynaklar: list[str]
    atlanan: list[str]


def bilgileri_coz(metin: str) -> dict:
    sonuc = {}
    for satir in metin.splitlines():
        if ":" not in satir:
            continue
        anahtar, deger = satir.split(":", 1)
        k = katla(anahtar)
        alan = next((a for a, adlar in ALANLAR.items() if k in adlar), None)
        if alan and deger.strip() and alan not in sonuc:
            sonuc[alan] = deger.strip()
    return sonuc


def talepleri_coz(metin: str) -> list[str]:
    parcalar = [re.sub(r"^\s*(?:\d+[.)-]|[-*•])\s*", "", s).strip() for s in metin.splitlines()]
    return [p for p in parcalar if p]


def delilleri_oku(yol: Path) -> list[Delil]:
    metin = belge.txt_oku(yol)
    ilk = "\n".join(metin.splitlines()[:5])
    satirlar = [r for r in csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)) if any(c.strip() for c in r)]
    if not satirlar:
        return []
    b = [katla(c) for c in satirlar[0]]
    j = lambda *adlar: next((b.index(a) for a in adlar if a in b), None)  # noqa: E731
    jn, ja, jt, je = j("delil no", "no"), j("delil", "aciklama", "delil adi"), j("tur", "delil turu"), j("ek", "ek no")
    if ja is None:
        raise ValueError(f"{yol.name}: 'Delil' sütunu bulunamadı")
    al = lambda r, i: r[i].strip() if i is not None and i < len(r) else ""  # noqa: E731
    return [Delil(al(r, jn) or f"D{n}", al(r, ja), al(r, jt), al(r, je)) for n, r in enumerate(satirlar[1:], 1) if al(r, ja)]


MADDE_RE = re.compile(r"^\s*(?:MADDE|Madde)\s+(\d+[A-Za-z/]*)\s*[-–—:.]?\s*(.*)$", re.M)


def maddelere_bol(metin: str, kid: str, kaynak: str) -> list[Madde]:
    es = list(MADDE_RE.finditer(metin))
    if not es:
        return [Madde(f"{kid}-tum", kaynak, "", metin.strip())]
    return [Madde(f"{kid}-md{m.group(1)}", kaynak, m.group(1), (m.group(2) + "\n" + metin[m.end():es[i + 1].start() if i + 1 < len(es) else len(metin)]).strip())
            for i, m in enumerate(es)]


def dosyalari_oku(klasor: Path) -> Dava:
    bilgi, olay, talepler, deliller, maddeler, kaynaklar, atlanan = {}, "", [], [], [], [], []
    for yol in sorted(p for p in klasor.rglob("*") if p.is_file()):
        ad = katla(yol.stem)
        mevzuat = "mevzuat" in katla(yol.parent.name) or ad.startswith(("kanun", "mevzuat"))
        try:
            if yol.suffix.lower() in {".csv"} and "delil" in ad:
                deliller += delilleri_oku(yol)
            elif yol.suffix.lower() not in belge.DESTEKLENEN:
                atlanan.append(yol.name)
            elif mevzuat:
                kid = f"M{len(kaynaklar) + 1}"
                metin = belge.metin_oku(yol)
                baslik = next((s.strip() for s in metin.splitlines() if s.strip()), yol.stem)[:90]
                kaynaklar.append(f"{kid}: {baslik} ({yol.name})")
                maddeler += maddelere_bol(metin, kid, baslik)
            elif "bilgi" in ad:
                bilgi.update(bilgileri_coz(belge.metin_oku(yol)))
            elif "olay" in ad or "ozet" in ad:
                olay = (olay + "\n\n" + belge.metin_oku(yol)).strip()
            elif "talep" in ad or "istem" in ad:
                talepler += talepleri_coz(belge.metin_oku(yol))
            else:
                atlanan.append(yol.name)
        except (belge.BelgeHatasi, ValueError, OSError) as h:
            atlanan.append(f"{yol.name} ({h})")
    return Dava(bilgi, olay, talepler, deliller, maddeler, kaynaklar, atlanan)


# ----------------------------------------------------------------------------
# Kod kontrolleri (HMK md. 119 kontrol listesi)
# ----------------------------------------------------------------------------

def kontrol_listesi(d: Dava) -> tuple[list[dict], list[tuple[str, str, str]]]:
    b = d.bilgi
    liste, sorunlar = [], []

    def madde(bent, unsur, durum, aciklama=""):
        liste.append({"bent": bent, "unsur": unsur, "durum": durum, "aciklama": aciklama})
        if durum == "Eksik":
            sorunlar.append(("yüksek", f"HMK 119/1-{bent}: {unsur} eksik" + (f" — {aciklama}" if aciklama else ""), unsur))
        elif durum == "Kontrol":
            sorunlar.append(("orta", f"HMK 119/1-{bent}: {unsur} — {aciklama}", unsur))

    madde("a", "Mahkemenin adı", "Var" if b.get("mahkeme") else "Eksik")
    for taraf in ("davaci", "davali"):
        ad = "Davacı" if taraf == "davaci" else "Davalı"
        madde("b", f"{ad} adı / unvanı", "Var" if b.get(taraf) else "Eksik")
        madde("b", f"{ad} adresi", "Var" if b.get(f"{taraf}_adres") else "Eksik", "tebligat yapılabilecek adres")
    davaci = b.get("davaci", "")
    if davaci and tuzel_mi(davaci):
        vkn = re.sub(r"\D", "", b.get("davaci_vkn", ""))
        if not vkn and not b.get("davaci_mersis"):
            madde("c", "Davacı kimlik / vergi numarası", "Kontrol", "tüzel kişi davacının vergi kimlik no veya MERSİS no'su yazılmamış")
        elif vkn and not vkn_gecerli(vkn):
            madde("c", "Davacı vergi kimlik no", "Kontrol", f"{vkn}: kontrol hanesi tutmuyor")
        else:
            madde("c", "Davacı vergi kimlik / MERSİS no (tüzel kişi)", "Var")
    else:
        tckn = re.sub(r"\D", "", b.get("davaci_tckn", ""))
        if not tckn:
            madde("c", "Davacı T.C. kimlik numarası", "Eksik", "gerçek kişi davacıda zorunlu")
        elif not tckn_gecerli(tckn):
            madde("c", "Davacı T.C. kimlik numarası", "Kontrol", "kontrol haneleri tutmuyor")
        else:
            madde("c", "Davacı T.C. kimlik numarası", "Var")
    vkn_d = re.sub(r"\D", "", b.get("davali_vkn", ""))
    if vkn_d and not vkn_gecerli(vkn_d):
        sorunlar.append(("orta", f"Davalı vergi kimlik no {vkn_d}: kontrol hanesi tutmuyor", "Davalı"))
    madde("ç", "Davacı vekilinin adı ve adresi", ("Var" if b.get("vekil_adres") else "Kontrol") if b.get("vekil") else "Yok (vekilsiz)",
          "" if not b.get("vekil") or b.get("vekil_adres") else "vekil adresi yazılmamış")
    madde("d", "Davanın konusu", "Var" if b.get("konu") else "Eksik")
    madde("d", "Dava değeri", "Var" if b.get("deger") else "Kontrol", "" if b.get("deger") else "malvarlığı haklarına ilişkin davalarda zorunlu")
    madde("e", "Vakıaların özeti (olay özeti)", "Var" if len(d.olay) > 50 else "Eksik")
    madde("f", "Deliller", "Var" if d.deliller else "Eksik")
    madde("g", "Hukuki sebepler", "Taslakta", "model yazar, avukat kontrol eder")
    madde("ğ", "Talep sonucu", "Var" if d.talepler else "Eksik")
    madde("h", "İmza", "Avukat", "taslak avukat tarafından kontrol edilip imzalanacak")

    tur = katla(b.get("tur", "") + " " + b.get("konu", ""))
    arabulucu_var = "arabulu" in katla(d.olay + " " + " ".join(x.ad for x in d.deliller))
    if any(k in tur for k in ARABULUCULUK_IPUCU):
        if arabulucu_var:
            sorunlar.append(("bilgi", "Dava türü zorunlu arabuluculuğa tabi olabilir; arabuluculuk son tutanağı dosyada var, aslı veya onaylı örneği "
                                      "dilekçeye eklenmeli", "Arabuluculuk"))
        else:
            sorunlar.append(("yüksek", "Dava türü zorunlu arabuluculuğa tabi olabilir (6325 sayılı Kanun md. 18/A ve özel kanunlar): arabuluculuk son "
                                       "tutanağı bulunamadı. Dava şartı olup olmadığını kontrol edin.", "Arabuluculuk"))
    sorunlar.append(("bilgi", "Harç, gider avansı ve görevli / yetkili mahkeme avukat tarafından kontrol edilmeli", "Usul"))
    return liste, sorunlar


# ----------------------------------------------------------------------------
# Maskeleme
# ----------------------------------------------------------------------------

def takma_adlar(d: Dava, terimler: list[str]) -> dict[str, str]:
    harita = {}
    for alan, takma in (("davaci", "[DAVACI]"), ("davali", "[DAVALI]"), ("vekil", "[DAVACI VEKİLİ]")):
        ad = d.bilgi.get(alan, "").strip()
        if ad:
            harita[ad] = takma
            sade = re.sub(r"^(Av\.|Avukat)\s*", "", ad).strip()
            if sade != ad:
                harita.setdefault(sade, takma)
    for i, ad in enumerate(dict.fromkeys(t.strip() for t in terimler if t.strip()), 1):
        harita.setdefault(ad, f"[KİŞİ-{i}]")
    return harita


def maskele(metin: str, harita: dict[str, str]) -> str:
    for gercek, takma in sorted(harita.items(), key=lambda x: -len(x[0])):
        metin = re.sub(re.escape(gercek), takma, metin, flags=re.I)
    return llm.maskele(metin)


def geri_ac(metin: str, harita: dict[str, str]) -> str:
    ters = {}
    for gercek, takma in harita.items():
        ters.setdefault(takma, gercek)
    for takma, gercek in sorted(ters.items(), key=lambda x: -len(x[0])):
        metin = metin.replace(takma, gercek)
    return metin


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

SEMA = {
    "type": "object",
    "properties": {
        "aciklamalar": {"type": "array", "items": {
            "type": "object",
            "properties": {"metin": {"type": "string"}, "dayanak_alinti": {"type": "string"}, "deliller": {"type": "array", "items": {"type": "string"}}},
            "required": ["metin", "dayanak_alinti", "deliller"], "additionalProperties": False}},
        "hukuki_sebepler": {"type": "array", "items": {
            "type": "object", "properties": {"ifade": {"type": "string"}, "madde": {"type": "string"}},
            "required": ["ifade", "madde"], "additionalProperties": False}},
        "sonuc_ve_istem": {"type": "array", "items": {"type": "string"}},
        "avukata_notlar": {"type": "array", "items": {"type": "string"}},
    },
    "required": ["aciklamalar", "hukuki_sebepler", "sonuc_ve_istem", "avukata_notlar"],
    "additionalProperties": False,
}


def duz(s: str) -> str:
    return " ".join(katla(s).split())


def sayilar(metin: str) -> set[str]:
    return {re.sub(r"\D", "", m) for m in re.findall(r"\d[\d.,]*\d", metin) if len(re.sub(r"\D", "", m)) >= 3}


def hazirla(d: Dava, kontroller, harita) -> tuple[str, str]:
    b = d.bilgi
    dava = "\n".join(f"{et}: {b[a]}" for et, a in (("Mahkeme", "mahkeme"), ("Davacı", "davaci"), ("Davacı vekili", "vekil"), ("Davalı", "davali"),
                                                   ("Dava türü", "tur"), ("Dava değeri", "deger"), ("Konu", "konu")) if b.get(a))
    mesaj = "\n".join([
        "<dava>", maskele(dava, harita), "</dava>",
        "<olay_ozeti>", maskele(d.olay, harita), "</olay_ozeti>",
        "<talepler>", *[f"{i}. {maskele(t, harita)}" for i, t in enumerate(d.talepler, 1)], "</talepler>",
        "<deliller>", *[f"{x.no}: {maskele(x.ad, harita)}" + (f" ({x.tur})" if x.tur else "") for x in d.deliller], "</deliller>",
        "<mevzuat>", *([f'<madde id="{m.id}" kaynak="{m.kaynak}" no="{m.no}">\n{m.metin}\n</madde>' for m in d.maddeler] or ["Verilmedi."]), "</mevzuat>",
        "<kanunlar>", *[f"{n} sayılı {ad}" for n, ad in KANUNLAR.items()], "</kanunlar>",
        "<kod_kontrolleri>", *([f"- [{o}] {a}" for o, a, _ in kontroller] or ["- Kontrol bulgusu yok"]), "</kod_kontrolleri>",
    ])
    return (BURASI / "prompt.md").read_text(encoding="utf-8"), mesaj


MADDE_ATIF = re.compile(r"(?:\bmd\.?|\bmadde(?:si|sinin|sine)?)\s*(\d+)|(\d+)\s*\.\s*madde", re.I)


def denetle(y: dict, d: Dava, mesaj: str) -> list[tuple[str, str, str]]:
    sorunlar = []
    olay = duz(d.olay)
    delil = {x.no: x for x in d.deliller}
    mad = {m.id: m for m in d.maddeler}
    for k, a in enumerate(y["aciklamalar"], 1):
        a["sira"] = k
        if not a["dayanak_alinti"].strip() or duz(a["dayanak_alinti"]) not in olay:
            a["dayanak"] = "Doğrulanamadı"
            sorunlar.append(("yüksek", f"Açıklama {k}: dayanak alıntısı olay özetinde birebir geçmiyor — vakıa avukatça kontrol edilmeli", f"Açıklama {k}"))
        else:
            a["dayanak"] = "Doğrulandı"
        bilinmeyen = [n for n in a["deliller"] if n not in delil]
        if bilinmeyen:
            sorunlar.append(("yüksek", f"Açıklama {k}: listede olmayan delil numarası yok sayıldı: {', '.join(bilinmeyen)}", f"Açıklama {k}"))
        a["deliller"] = [n for n in a["deliller"] if n in delil]
        for n in a["deliller"]:
            delil[n].vakialar.append(k)
        if not a["deliller"]:
            sorunlar.append(("orta", f"Açıklama {k}: vakıa için delil gösterilmemiş (HMK md. 119/1-f)", f"Açıklama {k}"))
    for x in d.deliller:
        if not x.vakialar:
            sorunlar.append(("bilgi", f"{x.no} ({x.ad}) hiçbir vakıada gösterilmemiş; deliller listesinde yer alır", x.no))
    verilen_no = {m.no for m in d.maddeler if m.no}
    for h in y["hukuki_sebepler"]:
        h["durum"] = "—"
        if h["madde"]:
            m = mad.get(h["madde"])
            if m is None:
                h["durum"] = "Madde verilmemiş"
                sorunlar.append(("yüksek", f"Hukuki sebep '{h['ifade'][:60]}': '{h['madde']}' verilen mevzuatta yok", "Hukuki sebepler"))
            else:
                h["durum"] = f"{m.kaynak} md. {m.no}" if m.no else m.kaynak
        atiflar = {a or b for a, b in MADDE_ATIF.findall(h["ifade"])}
        dogrulanamayan = sorted(atiflar - verilen_no, key=lambda s: int(s))
        if dogrulanamayan:
            h["durum"] = "Madde no doğrulanamadı"
            sorunlar.append(("yüksek", f"Hukuki sebep '{h['ifade'][:60]}': madde {', '.join(dogrulanamayan)} verilen mevzuatta yok — avukatça kontrol "
                                       "edilmeli", "Hukuki sebepler"))
        for n in re.findall(r"(\d+)\s*sayılı", h["ifade"]):
            if n not in KANUNLAR:
                sorunlar.append(("orta", f"Hukuki sebep: {n} sayılı kanun listede yok — kontrol edin", "Hukuki sebepler"))
    deger = sayilar(d.bilgi.get("deger", ""))
    sonuc = " ".join(y["sonuc_ve_istem"])
    if deger and not deger & sayilar(sonuc):
        sorunlar.append(("orta", f"Dava değeri ({d.bilgi['deger']}) sonuç ve istemde geçmiyor", "Sonuç ve istem"))
    if len(y["sonuc_ve_istem"]) < len(d.talepler):
        sorunlar.append(("orta", f"{len(d.talepler)} talep var, sonuç ve istemde {len(y['sonuc_ve_istem'])} madde — eksik talep olabilir", "Sonuç ve istem"))
    bilinen = sayilar(mesaj) | {x for x in re.findall(r"\d+", mesaj) if len(x) >= 3}
    metin = " ".join([*(a["metin"] for a in y["aciklamalar"]), *(h["ifade"] for h in y["hukuki_sebepler"]), sonuc])
    y["dogrulanamayan_sayilar"] = sorted(s for s in sayilar(metin) if s not in bilinen)
    if y["dogrulanamayan_sayilar"]:
        sorunlar.append(("yüksek", f"Girdilerde olmayan sayı / tarih: {', '.join(y['dogrulanamayan_sayilar'])}", "Sayılar"))
    return sorunlar


def model_yaz(d: Dava, kontroller, harita) -> dict:
    sistem, mesaj = hazirla(d, kontroller, harita)
    y = llm.json_iste(sistem, mesaj, SEMA)
    y["sorunlar"] = denetle(y, d, mesaj)
    return y


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
ONEM = {"yüksek": "FDE2E1", "orta": "FFF4CE", "bilgi": "E8F0FE"}
DURUM = {"Var": "E3F4E1", "Eksik": "FDE2E1", "Kontrol": "FFF4CE", "Doğrulandı": "E3F4E1", "Doğrulanamadı": "FDE2E1"}
KARAR = PatternFill("solid", fgColor="FFF4CE")
MODEL = PatternFill("solid", fgColor="EDE7FF")
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def dilekce_md(d: Dava, y: dict, harita: dict, tum_sorunlar) -> str:
    g = lambda s: geri_ac(s, harita)  # noqa: E731
    b = d.bilgi
    s = ["> **TASLAKTIR — avukat tarafından kontrol edilip imzalanmadan kullanılamaz.**", "",
         f"**{buyuk(b.get('mahkeme', '[MAHKEME ADI]'))} SAYIN HÂKİMLİĞİNE**", ""]
    davaci = b.get("davaci", "[DAVACI]")
    kimlik = (f"(VKN: {b['davaci_vkn']})" if b.get("davaci_vkn") else f"(MERSİS: {b['davaci_mersis']})" if b.get("davaci_mersis")
              else f"(TCKN: {b['davaci_tckn']})" if b.get("davaci_tckn") else "(TCKN / VKN: …)")
    s += [f"**DAVACI** : {davaci} {kimlik}  ", f"{'':14}{b.get('davaci_adres', '[ADRES]')}  "]
    if b.get("vekil"):
        s += [f"**VEKİLİ** : {b['vekil']}  ", f"{'':14}{b.get('vekil_adres', '[ADRES]')}  "]
    davali_k = f" (VKN: {b['davali_vkn']})" if b.get("davali_vkn") else ""
    s += [f"**DAVALI** : {b.get('davali', '[DAVALI]')}{davali_k}  ", f"{'':14}{b.get('davali_adres', '[ADRES]')}  "]
    if b.get("deger"):
        s += [f"**DAVA DEĞERİ** : {b['deger']}  "]
    s += [f"**KONU** : {b.get('konu', '[KONU]')}", "", "**AÇIKLAMALAR** :", ""]
    for a in y["aciklamalar"]:
        delil = f" (Delil: {', '.join(a['deliller'])})" if a["deliller"] else " **[delil gösterilmedi]**"
        isaret = "" if a["dayanak"] == "Doğrulandı" else " **[dayanak doğrulanamadı]**"
        s += [f"{a['sira']}. {g(a['metin'])}{delil}{isaret}", ""]
    s += ["**HUKUKİ SEBEPLER** :", ""]
    s += [f"- {g(h['ifade'])}" + (f" **[{h['durum']}]**" if h["durum"] in ("Madde verilmemiş", "Madde no doğrulanamadı") else "")
          for h in y["hukuki_sebepler"]] or ["- [Avukat tarafından eklenecek]"]
    s += ["", "**HUKUKİ DELİLLER** :", ""]
    s += [f"{i}. {x.ad}" + (f" ({x.ek})" if x.ek else "") for i, x in enumerate(d.deliller, 1)]
    s += ["", "ve sair her türlü yasal delil.", "", "**SONUÇ VE İSTEM** : Yukarıda açıklanan ve re'sen dikkate alınacak nedenlerle;", ""]
    s += [f"{i}. {g(t)}" for i, t in enumerate(y["sonuc_ve_istem"], 1)]
    s += ["", "karar verilmesini saygılarımızla vekaleten arz ve talep ederiz." if b.get("vekil") else "karar verilmesini saygılarımla arz ve talep ederim.",
          "", f"{'':50}{(b.get('vekil') or davaci)}", f"{'':50}(imza)", f"{'':50}… / … / {date.today().year}", ""]
    ekler = [x for x in d.deliller if x.ek]
    if ekler or b.get("vekil"):
        s += ["**EKLER** :", ""]
        s += [f"- {x.ek}: {x.ad}" for x in ekler]
        if b.get("vekil"):
            s += ["- Vekâletname örneği"]
        s += [""]
    s += ["---", "", "## Avukata notlar (dilekçeye dahil değildir)", ""]
    s += [f"- {g(n)}" for n in y["avukata_notlar"]] or ["- —"]
    if tum_sorunlar:
        s += ["", "## Kod kontrolleri", ""] + [f"- **{o}** {g(a)}" for o, a, _ in tum_sorunlar if o != "bilgi"]
        s += [f"- {g(a)}" for o, a, _ in tum_sorunlar if o == "bilgi"]
    return "\n".join(s) + "\n"


def rapor_yaz(cikti: Path, d: Dava, liste, kontroller, y, harita) -> Path:
    g = lambda s: geri_ac(s, harita)  # noqa: E731
    tum = kontroller + y["sorunlar"]
    wb = Workbook()
    kl = wb.active
    kl.title = "Kontrol Listesi"
    _baslik(kl, ["HMK 119/1", "Unsur", "Durum", "Açıklama", "Avukat Kontrolü"], (9, 40, 12, 60, 18))
    for x in liste:
        kl.append([x["bent"], x["unsur"], x["durum"], x["aciklama"], ""])
        if x["durum"] in DURUM:
            kl.cell(kl.max_row, 3).fill = PatternFill("solid", fgColor=DURUM[x["durum"]])
        kl.cell(kl.max_row, 5).fill = KARAR

    vd = wb.create_sheet("Vakıa–Delil")
    _baslik(vd, ["No", "Açıklama (taslak)", "Dayanak Alıntısı (olay özeti)", "Dayanak", "Deliller", "Avukat Kontrolü"], (5, 70, 50, 14, 14, 18))
    for a in y["aciklamalar"]:
        vd.append([a["sira"], g(a["metin"]), g(a["dayanak_alinti"]), a["dayanak"], ", ".join(a["deliller"]) or "YOK", ""])
        vd.cell(vd.max_row, 2).fill = MODEL
        vd.cell(vd.max_row, 4).fill = PatternFill("solid", fgColor=DURUM[a["dayanak"]])
        if not a["deliller"]:
            vd.cell(vd.max_row, 5).fill = PatternFill("solid", fgColor=ONEM["orta"])
        vd.cell(vd.max_row, 6).fill = KARAR
        for h in vd[vd.max_row]:
            h.alignment = UST
    vd.append([])
    _baslik(vd, ["Delil", "Açıklama", "Tür", "Ek", "Vakıalar", ""], ())
    for x in d.deliller:
        vd.append([x.no, x.ad, x.tur, x.ek, ", ".join(map(str, x.vakialar)) or "—"])

    hs = wb.create_sheet("Hukuki Sebepler")
    _baslik(hs, ["İfade (taslak)", "Dayanak / Durum", "Avukat Kontrolü"], (90, 30, 18))
    for h in y["hukuki_sebepler"]:
        hs.append([g(h["ifade"]), h["durum"], ""])
        hs.cell(hs.max_row, 1).fill = MODEL
        if "doğrulanamadı" in h["durum"] or "verilmemiş" in h["durum"]:
            hs.cell(hs.max_row, 2).fill = PatternFill("solid", fgColor=ONEM["yüksek"])
        hs.cell(hs.max_row, 3).fill = KARAR
        hs.cell(hs.max_row, 1).alignment = UST
    hs.append([])
    hs.append(["Verilen mevzuat: " + ("; ".join(d.kaynaklar) or "yok — madde numaraları avukat tarafından eklenmeli")])

    kt = wb.create_sheet("Kontroller")
    _baslik(kt, ["Önem", "Kontrol", "İlgili", "İnceleme"], (9, 100, 18, 22))
    for on, a, il in tum:
        kt.append([on, g(a), il, ""])
        kt.cell(kt.max_row, 1).fill = PatternFill("solid", fgColor=ONEM[on])
        kt.cell(kt.max_row, 2).alignment = UST
        kt.cell(kt.max_row, 4).fill = KARAR

    an = wb.create_sheet("Avukata Notlar")
    _baslik(an, ["No", "Not (model)", "Değerlendirme"], (5, 100, 30))
    for i, n in enumerate(y["avukata_notlar"], 1):
        an.append([i, g(n), ""])
        an.cell(an.max_row, 2).fill = MODEL
        an.cell(an.max_row, 2).alignment = UST
        an.cell(an.max_row, 3).fill = KARAR
    an.append([])
    an.append(["", f"Model: {llm.kullanim_ozeti()} · okunamayan dosyalar: {'; '.join(d.atlanan) or '—'}"])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    md = cikti.with_suffix(".md")
    md.write_text(dilekce_md(d, y, harita, tum), encoding="utf-8")
    return md


def calistir(klasor: Path, cikti: Path, terimler: list[str] | None = None, evet: bool = False) -> dict:
    if not klasor.is_dir():
        raise llm.LLMHatasi(f"{klasor} klasörü bulunamadı.")
    d = dosyalari_oku(klasor)
    if not d.olay:
        raise llm.LLMHatasi(f"{klasor}: olay özeti bulunamadı (adında 'olay' veya 'ozet' geçen .txt/.docx/.pdf).")
    liste, kontroller = kontrol_listesi(d)
    harita = takma_adlar(d, terimler or [])
    print(f"[OK] Olay özeti {len(d.olay)} karakter · {len(d.talepler)} talep · {len(d.deliller)} delil · {len(d.maddeler)} mevzuat maddesi")
    llm.onay_al(f"Olay özeti, {len(d.talepler)} talep, {len(d.deliller)} delil ve {len(d.maddeler)} mevzuat maddesi gönderilecek (taraf adları takma "
                "adlı; adres ve kimlik numaraları gönderilmez; telefon, e-posta, TCKN, IBAN maskeli).", evet)
    y = model_yaz(d, kontroller, harita)
    md = rapor_yaz(cikti, d, liste, kontroller, y, harita)
    return {"dava": d, "liste": liste, "kontroller": kontroller, "yanit": y, "harita": harita, "md": md}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Olay özeti, talepler ve deliller listesinden dava dilekçesi taslağı hazırlar (avukat kontrolü zorunludur).")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "ornek_alacak_davasi",
                   help="Klasör: dava_bilgileri.txt, olay_ozeti.txt, talepler.txt, deliller.csv, [mevzuat/]")
    p.add_argument("--gizle", nargs="*", default=[], metavar="AD", help="Ayrıca maskelenecek adlar (tanık, yetkili vb.)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "dilekce_taslagi.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.cikti, a.gizle, a.evet)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    for o, ac, _ in s["kontroller"] + s["yanit"]["sorunlar"]:
        if o != "bilgi":
            print(f"[{'X' if o == 'yüksek' else '!'}] {ac}")
    print(f"[OK] {len(s['yanit']['aciklamalar'])} vakıa · dilekçe taslağı: {s['md'].resolve()} · kontroller: {a.cikti.resolve()}")
    print("[i] Taslaktır; avukat kontrolü ve imzası olmadan kullanılamaz.")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
