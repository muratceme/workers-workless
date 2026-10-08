"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

DOSYA = KLASOR / "ornek_veri" / "ornek_alacak_davasi"
MESAJLAR = []


def aciklama(metin, alinti, deliller):
    return {"metin": metin, "dayanak_alinti": alinti, "deliller": deliller}


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return {
        "aciklamalar": [
            aciklama("[DAVACI] ile [DAVALI] arasında 05.01.2026 tarihli çerçeve satış sözleşmesi imzalanmıştır.",
                     "arasında 05.01.2026 tarihli çerçeve satış sözleşmesi imzalanmıştır", ["D1"]),
            aciklama("Müvekkil şirket malları teslim etmiş ve toplam 412.500 TL tutarında üç fatura düzenlemiştir.",
                     "toplam 412.500 TL tutarında üç fatura düzenlenmiştir", ["D2", "D3"]),
            aciklama("Davalı 15.04.2026 tarihinde 100.000 TL ödemiş, bakiye 312.500 TL ödenmemiştir.",
                     "Davalı 15.04.2026 tarihinde 100.000 TL ödeme yapmış, bakiye 312.500 TL ödenmemiştir", ["D4", "D7", "D99"]),
            aciklama("İhtarname 12.06.2026 tarihinde tebliğ edilmiş, süre 17.06.2026 tarihinde dolmuştur.",
                     "İhtarname 12.06.2026 tarihinde tebliğ edilmiş", ["D5"]),
            aciklama("Arabuluculuk görüşmelerinde taraflar anlaşamamıştır.", "20.08.2026 tarihli son tutanakla taraflar anlaşamamıştır", ["D6"]),
            aciklama("Davalı malları eksiksiz kabul etmiş, 450.000 TL'lik ek sipariş vermiştir.", "Davalı malları eksiksiz kabul etmiştir", []),
        ],
        "hukuki_sebepler": [
            {"ifade": "6098 sayılı Türk Borçlar Kanunu, 6102 sayılı Türk Ticaret Kanunu ve ilgili mevzuat", "madde": ""},
            {"ifade": "6100 sayılı Hukuk Muhakemeleri Kanunu md. 119", "madde": ""},
            {"ifade": "9999 sayılı Uydurma Kanun", "madde": ""},
        ],
        "sonuc_ve_istem": ["312.500 TL alacağın 17.06.2026 temerrüt tarihinden itibaren işleyecek avans faiziyle birlikte davalıdan tahsiline,",
                           "Yargılama giderleri ile vekâlet ücretinin davalıya yükletilmesine,"],
        "avukata_notlar": ["[DAVALI] ayıp savunmasında bulunabilir; ayıp ihbarı süreleri kontrol edilmeli."],
    }


class GirdiTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.d = agent.dosyalari_oku(DOSYA)

    def test_okuma(self):
        b = self.d.bilgi
        self.assertEqual((b["mahkeme"], b["davaci_vkn"], b["deger"]), ("İstanbul Anadolu Nöbetçi Asliye Ticaret Mahkemesi", "1234567890", "312.500 TL"))
        self.assertEqual(len(self.d.talepler), 2)
        self.assertTrue(self.d.talepler[0].startswith("312.500 TL"))
        self.assertEqual([x.no for x in self.d.deliller], [f"D{i}" for i in range(1, 10)])
        self.assertEqual(self.d.deliller[5].ek, "Ek-6")
        self.assertEqual(self.d.maddeler, [])

    def test_kontrol_listesi(self):
        liste, k = agent.kontrol_listesi(self.d)
        durum = {(x["bent"], x["unsur"]): x["durum"] for x in liste}
        self.assertEqual(durum[("c", "Davacı vergi kimlik / MERSİS no (tüzel kişi)")], "Var")
        self.assertNotIn("Eksik", durum.values())
        self.assertTrue(any(o == "bilgi" and "arabuluculuk son tutanağı dosyada var" in a for o, a, _ in k))

    def test_eksik_gercek_kisi_ve_arabuluculuk(self):
        d = agent.Dava({"davaci": "Ali Deneme", "davaci_tckn": "12345678901", "tur": "Kira alacağı"}, "Kısa", [], [], [], [], [])
        liste, k = agent.kontrol_listesi(d)
        durum = {x["unsur"]: x["durum"] for x in liste}
        self.assertEqual(durum["Davacı T.C. kimlik numarası"], "Kontrol")
        self.assertEqual(durum["Mahkemenin adı"], "Eksik")
        self.assertTrue(any(o == "yüksek" and "arabuluculuk son tutanağı bulunamadı" in a for o, a, _ in k))
        self.assertTrue(agent.tckn_gecerli("10000000146"))
        self.assertTrue(agent.vkn_gecerli("2468135797"))


class AgentTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "d.xlsx"
        MESAJLAR.clear()
        with mock.patch.object(agent.llm, "json_iste", sahte_json_iste):
            cls.s = agent.calistir(DOSYA, cls.cikti, ["Sevkiyat Sorumlusu"], evet=True)
        cls.y = cls.s["yanit"]
        cls.sorun = [a for _, a, _ in cls.y["sorunlar"]]

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_maskeleme(self):
        m = MESAJLAR[0]
        self.assertIn("[DAVACI]", m)
        self.assertIn("[DAVALI]", m)
        for gizli in ("Örnek Tekstil", "Kurgu Mağazacılık", "Deneme Avukat", "1234567890", "Kadıköy / İstanbul"):
            self.assertNotIn(gizli, m)

    def test_denetim(self):
        self.assertEqual([a["dayanak"] for a in self.y["aciklamalar"]], ["Doğrulandı"] * 5 + ["Doğrulanamadı"])
        self.assertTrue(any("Açıklama 6: vakıa için delil gösterilmemiş" in a for a in self.sorun))
        self.assertTrue(any("D99" in a for a in self.sorun))
        self.assertEqual(self.y["aciklamalar"][2]["deliller"], ["D4", "D7"])
        self.assertIn("450000", self.y["dogrulanamayan_sayilar"])
        self.assertTrue(any("madde 119 verilen mevzuatta yok" in a for a in self.sorun))
        self.assertTrue(any("9999 sayılı kanun listede yok" in a for a in self.sorun))
        self.assertTrue(any("D8 (Bilirkişi incelemesi) hiçbir vakıada" in a for a in self.sorun))
        self.assertFalse(any("Dava değeri" in a for a in self.sorun))

    def test_dilekce(self):
        md = self.s["md"].read_text(encoding="utf-8")
        self.assertIn("İSTANBUL ANADOLU NÖBETÇİ ASLİYE TİCARET MAHKEMESİ SAYIN HÂKİMLİĞİNE", md)
        self.assertIn("**DAVACI** : Örnek Tekstil San. ve Tic. A.Ş. (VKN: 1234567890)", md)
        self.assertIn("1. Örnek Tekstil San. ve Tic. A.Ş. ile Kurgu Mağazacılık Ltd. Şti. arasında", md)
        self.assertIn("[dayanak doğrulanamadı]", md)
        self.assertIn("md. 119 **[Madde no doğrulanamadı]**", md)
        self.assertIn("- Ek-6: Arabuluculuk son tutanağı (20.08.2026)", md)
        self.assertIn("TASLAKTIR", md)

    def test_excel(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Kontrol Listesi", "Vakıa–Delil", "Hukuki Sebepler", "Kontroller", "Avukata Notlar"])
        self.assertEqual(wb["Vakıa–Delil"]["E7"].value, "YOK")
        self.assertIn("Kurgu Mağazacılık", wb["Avukata Notlar"]["B2"].value)


class CliTesti(unittest.TestCase):
    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(agent.main(["--girdi", str(Path(tmp) / "yok")]), 1)
            self.assertEqual(agent.main(["--girdi", tmp]), 1)                # olay özeti yok
            with mock.patch.object(agent.llm, "json_iste", sahte_json_iste):
                self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "x.xlsx")]), 0)


if __name__ == "__main__":
    unittest.main()
