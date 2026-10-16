"""
BT Olay Sonrası Raporu — Workers / Workless AI Agent
Bilgi Teknolojileri › Sistem Uzmanı

Olay sırasında tutulan zaman damgalı notlardan (sohbet dökümü, çağrı notları) suçlamasız olay sonrası raporu
(postmortem) hazırlar.
1. Kod notları okur: "Olay: No | Hizmet: … | Tarih: GG.AA.YYYY" başlığı, "SS:DD konuşan: mesaj" satırları ve "Aksiyon:"
   satırları. Gece yarısını geçen saatler ertesi güne yazılır. "Ad Soyad:" biçimindeki konuşanlar ve --gizle adları
   takma adla maskelenir.
2. Model özet, etki, zaman çizelgesi, kök neden (Doğrulandı / Varsayım), katkıda bulunan etkenler, aksiyonlar ve açık
   soruları yazar.
3. Kod denetler: zaman çizelgesindeki her damga notlarda aynen var mı ve sıralı mı; etki başlangıcı / tespit / hizmet
   düzeldi birer kez ve doğru sırada mı; sayı ve tarihler notlarda geçiyor mu; kişi adı var mı; "Doğrulandı" için notta
   teyit var mı; notlardaki her "Aksiyon:" satırı rapora alınmış mı, notta olmayan aksiyon eklenmiş mi. Sorun varsa
   model düzeltir.
4. Kod tespit, müdahale ve kesinti sürelerini hesaplar, --sla-dk hedefiyle karşılaştırır.

Kullanım:
    python agent.py                                                       # örnek olay
    python agent.py --notlar olay.txt --sla-dk 60 --gizle "Ad Soyad"
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}
TURLER = ["İlk belirti", "Etki başlangıcı", "Tespit", "Teşhis", "Geçici çözüm", "Hizmet düzeldi", "Kalıcı çözüm", "Bilgi"]
DOGRU = re.compile(r"\b(kok neden (dogrulandi|teyit edildi|kesinlesti)|dogrulandi|teyit edildi)\b")
SATIR = re.compile(r"^\s*(?:(\d{1,2}\.\d{1,2}\.\d{4})\s+)?(\d{1,2})[:.](\d{2})\s+(.*)$")
DAMGA = "%d.%m.%Y %H:%M"


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def sayilar(metin: str) -> set[str]:
    temiz = re.sub(r"\b\d{1,2}\.\d{1,2}\.\d{4}\b|\b[A-Z]{2,5}-[\w-]+\b", " ", str(metin))
    return {str(int(x)) if x.isdigit() else x.replace(",", ".") for x in re.findall(r"\d+(?:[.,]\d+)?", temiz)}


def tarihler(metin: str) -> set[str]:
    return set(re.findall(r"\b\d{2}\.\d{2}\.\d{4}\b", str(metin)))


def kelimeler(s: str) -> set[str]:
    return {w for w in katla(s).split() if len(w) >= 4}


def dk(a: datetime | None, b: datetime | None) -> int | None:
    return int((b - a).total_seconds() // 60) if a and b else None


# ----------------------------------------------------------------------------
# Girdi
# ----------------------------------------------------------------------------

def kisi_mi(ad: str) -> bool:
    k = ad.split()
    return 2 <= len(k) <= 4 and all(w[:1].isupper() for w in k) and not re.search(r"\d", ad)


def takma_adlar(adlar: list[str]) -> dict[str, str]:
    return {a: f"[KİŞİ-{i}]" for i, a in enumerate(dict.fromkeys(x.strip() for x in adlar if x.strip()), 1)}


def maskele(metin: str, kisiler: dict[str, str]) -> str:
    for g, t in sorted(kisiler.items(), key=lambda x: -len(x[0])):
        metin = re.sub(re.escape(g), t, metin, flags=re.I)
    return llm.maskele(metin)


def notlari_oku(yol: Path, gizle: list[str] | None = None) -> dict:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: not dosyası bulunamadı.")
    satirlar = belge.metin_oku(yol).replace("\r", "").splitlines()
    olay = {"no": yol.stem, "hizmet": "", "tarih": None}
    ham, aksiyon = [], []
    gun, onceki = None, None
    for s in satirlar:
        if not s.strip():
            continue
        m = re.match(r"^\s*Olay\s*:\s*(.+)$", s, re.I)
        if m:
            alan = [x.strip() for x in m[1].split("|")]
            olay["no"] = alan[0]
            for a in alan[1:]:
                if ":" in a:
                    k, v = (x.strip() for x in a.split(":", 1))
                    if katla(k) == "hizmet":
                        olay["hizmet"] = v
                    elif katla(k) == "tarih":
                        try:
                            gun = olay["tarih"] = datetime.strptime(v, "%d.%m.%Y")
                        except ValueError:
                            raise llm.LLMHatasi(f"{yol.name}: başlıktaki tarih okunamadı: '{v}'") from None
            continue
        m = re.match(r"^\s*Aksiyon\s*:\s*(.+)$", s, re.I)
        if m:
            aksiyon.append(m[1].strip())
            continue
        m = SATIR.match(s)
        if m:
            if m[1]:
                gun = datetime.strptime(m[1], "%d.%m.%Y")
            if gun is None:
                raise llm.LLMHatasi(f"{yol.name}: saat satırından önce tarih yok; başlığa 'Tarih: GG.AA.YYYY' ekleyin.")
            z = gun.replace(hour=int(m[2]), minute=int(m[3]))
            if not m[1] and onceki and z < onceki:                      # gece yarısını geçti
                gun += timedelta(days=1)
                z += timedelta(days=1)
            onceki = z
            kim, _, msj = m[4].partition(":") if re.match(r"^[^:]{1,40}:\s", m[4]) else ("", "", m[4])
            ham.append({"zaman": z, "kim": kim.strip(), "mesaj": msj.strip()})
        elif ham:
            ham[-1]["mesaj"] += " " + s.strip()
    if not ham:
        raise llm.LLMHatasi(f"{yol.name}: 'SS:DD konuşan: mesaj' biçiminde zaman damgalı satır yok.")
    kisiler = takma_adlar([h["kim"] for h in ham if kisi_mi(h["kim"])] + list(gizle or []))
    for h in ham:
        h["kim_m"] = maskele(h["kim"], kisiler) if h["kim"] else ""
        h["mesaj_m"] = maskele(h["mesaj"], kisiler)
        h["damga"] = h["zaman"].strftime(DAMGA)
    olay.update(satirlar=ham, aksiyonlar=[maskele(a, kisiler) for a in aksiyon], kisiler=kisiler)
    olay["metin"] = "\n".join(f"[{h['damga']}] {h['kim_m'] + ': ' if h['kim_m'] else ''}{h['mesaj_m']}" for h in ham)
    olay["metin"] += "".join(f"\nAksiyon: {a}" for a in olay["aksiyonlar"])
    return olay


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

SEMA = {"type": "object", "properties": {
    "baslik": S, "ozet": S, "etki": S, "etkilenen_hizmetler": {"type": "array", "items": S},
    "zaman_cizelgesi": {"type": "array", "items": {"type": "object", "properties": {"zaman": S, "tur": {"type": "string", "enum": TURLER}, "olay": S},
                                                   "required": ["zaman", "tur", "olay"], "additionalProperties": False}},
    "kok_neden": S, "kok_neden_durumu": {"type": "string", "enum": ["Doğrulandı", "Varsayım"]},
    "katkida_bulunanlar": {"type": "array", "items": S}, "iyi_gidenler": {"type": "array", "items": S},
    "aksiyonlar": {"type": "array", "items": {"type": "object", "properties": {"aksiyon": S, "tur": {"type": "string", "enum": ["Önleyici", "Tespit", "Süreç"]},
                                                                             "sahip": S, "termin": S},
                                              "required": ["aksiyon", "tur", "sahip", "termin"], "additionalProperties": False}},
    "acik_sorular": {"type": "array", "items": S}},
    "required": ["baslik", "ozet", "etki", "etkilenen_hizmetler", "zaman_cizelgesi", "kok_neden", "kok_neden_durumu", "katkida_bulunanlar", "iyi_gidenler",
                 "aksiyonlar", "acik_sorular"],
    "additionalProperties": False}


def benzer(a: str, b: str) -> float:
    ka, kb = kelimeler(a), kelimeler(b)
    return len(ka & kb) / min(len(ka), len(kb)) if ka and kb else 0.0


def anlar(c: dict) -> dict[str, datetime]:
    out = {}
    for x in c["zaman_cizelgesi"]:
        try:
            z = datetime.strptime(x["zaman"].strip(), DAMGA)
        except ValueError:
            continue
        out.setdefault(x["tur"], z)
    return out


def denetle(o: dict, c: dict) -> list[str]:
    sorun = []
    damgalar = {h["damga"] for h in o["satirlar"]}
    yanlis = [x["zaman"] for x in c["zaman_cizelgesi"] if x["zaman"].strip() not in damgalar]
    if yanlis:
        sorun.append(f"zaman çizelgesinde notlarda olmayan damga: {', '.join(yanlis)} (notlardaki [GG.AA.YYYY SS:DD] damgasını aynen yazın)")
    z = [datetime.strptime(x["zaman"].strip(), DAMGA) for x in c["zaman_cizelgesi"] if x["zaman"].strip() in damgalar]
    if any(b < a for a, b in zip(z, z[1:])):
        sorun.append("zaman çizelgesi sıralı değil")
    say = Counter(x["tur"] for x in c["zaman_cizelgesi"])
    for t in ("Etki başlangıcı", "Tespit"):
        if say[t] != 1:
            sorun.append(f"'{t}' zaman çizelgesinde {say[t]} kez var; tam bir kez olmalı")
    if say["Hizmet düzeldi"] > 1:
        sorun.append("'Hizmet düzeldi' birden çok kez var")
    a = anlar(c)
    for x, y in (("İlk belirti", "Tespit"), ("Etki başlangıcı", "Hizmet düzeldi"), ("Tespit", "Hizmet düzeldi"), ("Hizmet düzeldi", "Kalıcı çözüm")):
        if x in a and y in a and a[y] < a[x]:
            sorun.append(f"'{y}' ({a[y]:%H:%M}) '{x}' ({a[x]:%H:%M}) anından önce olamaz")
    metin = " ".join([c["baslik"], c["ozet"], c["etki"], c["kok_neden"], *c["katkida_bulunanlar"], *c["iyi_gidenler"], *c["acik_sorular"],
                      *(x["olay"] for x in c["zaman_cizelgesi"]), *(x["aksiyon"] + " " + x["sahip"] for x in c["aksiyonlar"])])
    fazla = sorted(sayilar(metin) - sayilar(o["metin"] + " " + o["no"]), key=lambda s: float(s))
    if fazla:
        sorun.append(f"notlarda olmayan sayı: {', '.join(fazla)} (süreleri kod hesaplar)")
    fazla_t = sorted(tarihler(metin + " " + " ".join(x["termin"] for x in c["aksiyonlar"])) - tarihler(o["metin"]))
    if fazla_t:
        sorun.append(f"notlarda olmayan tarih: {', '.join(fazla_t)}")
    kisi = sorted(set(re.findall(r"\[KİŞİ-\d+\]", metin)))
    if kisi:
        sorun.append(f"raporda kişi takma adı var ({', '.join(kisi)}); suçlamasız yazın, rol / ekip kullanın")
    if c["kok_neden_durumu"] == "Doğrulandı" and not DOGRU.search(katla(o["metin"])):
        sorun.append("notlarda kök nedenin doğrulandığı yazmıyor; durum 'Varsayım' olmalı")
    for n in o["aksiyonlar"]:
        if not any(benzer(n, x["aksiyon"]) >= 0.5 for x in c["aksiyonlar"]):
            sorun.append(f"notlardaki aksiyon rapora alınmamış: \"{n[:100]}\"")
    for x in c["aksiyonlar"]:
        if not any(benzer(n, x["aksiyon"]) >= 0.5 for n in o["aksiyonlar"]):
            sorun.append(f"notlarda olmayan aksiyon: \"{x['aksiyon'][:100]}\" (öneriyi açık sorulara yazın)")
        if not x["sahip"].strip():
            sorun.append(f"aksiyonun sahibi boş: \"{x['aksiyon'][:60]}\"")
    if not c["kok_neden"].strip():
        sorun.append("kök neden boş; bilinmiyorsa 'Varsayım' olarak en olası nedeni ve açık soruyu yazın")
    return sorun


def uret(o: dict, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    mesaj = f"<olay no=\"{o['no']}\" hizmet=\"{o['hizmet']}\">\n{o['metin']}\n</olay>"
    if duzeltme:
        mesaj += ("\n<duzeltme>\nÖnceki cevap: " + json.dumps(onceki, ensure_ascii=False) + "\nKodun bulduğu sorunlar:\n"
                  + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")
    return llm.json_iste(sistem, mesaj, SEMA)


def duzeltme_dongusu(uret_f, denetle_f, tur: int) -> tuple[dict, list[str], int]:
    taslak = uret_f(None, None)
    kalan = denetle_f(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret_f(kalan, taslak)
        kalan = denetle_f(taslak)
    return taslak, kalan, yapilan


# ----------------------------------------------------------------------------
# Hesap ve çıktı
# ----------------------------------------------------------------------------

def metrikler(c: dict, sla_dk: int | None) -> list[tuple[str, int | None, str]]:
    a = anlar(c)
    out = [("İlk belirtiden tespite (dk)", dk(a.get("İlk belirti"), a.get("Tespit")), "Uyarı / belirti ile BT'nin fark etmesi arası"),
           ("Etki başlangıcından tespite (dk)", dk(a.get("Etki başlangıcı"), a.get("Tespit")), ""),
           ("Tespitten hizmetin düzelmesine (dk)", dk(a.get("Tespit"), a.get("Hizmet düzeldi")), "Müdahale süresi"),
           ("Kesinti süresi (dk)", dk(a.get("Etki başlangıcı"), a.get("Hizmet düzeldi")), "Etki başlangıcı → hizmet düzeldi"),
           ("Kalıcı çözüme kadar (dk)", dk(a.get("Etki başlangıcı"), a.get("Kalıcı çözüm")), "")]
    k = out[3][1]
    if sla_dk is not None:
        out.append(("SLA hedefi (dk)", sla_dk, "Aşıldı" if k is not None and k > sla_dk else ("Karşılandı" if k is not None else "Kesinti süresi hesaplanamadı")))
    if "Hizmet düzeldi" not in a:
        out.append(("Durum", None, "Hizmetin düzeldiği an notlarda yok; olay açık ya da not eksik"))
    return out


def rapor_metni(o: dict, c: dict, m: list) -> str:
    l = [f"# Olay Sonrası Raporu · {o['no']}", "", f"**{c['baslik']}**", "",
         f"- Hizmet: {o['hizmet'] or '-'}", f"- Kök neden durumu: {c['kok_neden_durumu']}", ""]
    l += ["## Özet", "", c["ozet"], "", "## Etki", "", c["etki"]]
    if c["etkilenen_hizmetler"]:
        l += ["", "Etkilenen hizmetler: " + ", ".join(c["etkilenen_hizmetler"])]
    l += ["", "## Süreler (kod hesabı)", ""] + [f"- {a}: {'-' if v is None else v}{(' · ' + n) if n else ''}" for a, v, n in m]
    l += ["", "## Zaman çizelgesi", "", "| Zaman | Tür | Olay |", "|---|---|---|"]
    l += [f"| {x['zaman']} | {x['tur']} | {x['olay']} |" for x in c["zaman_cizelgesi"]]
    l += ["", "## Kök neden", "", c["kok_neden"], "", "## Katkıda bulunan etkenler", ""] + [f"- {x}" for x in c["katkida_bulunanlar"]]
    if c["iyi_gidenler"]:
        l += ["", "## İyi giden", ""] + [f"- {x}" for x in c["iyi_gidenler"]]
    l += ["", "## Aksiyonlar", "", "| Aksiyon | Tür | Sahip | Termin |", "|---|---|---|---|"]
    l += [f"| {x['aksiyon']} | {x['tur']} | {x['sahip']} | {x['termin'] or 'Belirlenecek'} |" for x in c["aksiyonlar"]]
    if c["acik_sorular"]:
        l += ["", "## Açık sorular", ""] + [f"- {x}" for x in c["acik_sorular"]]
    l += ["", "---", "Bu rapor suçlamasız yazılmıştır: amaç kişiyi değil süreç ve sistem açıklarını bulmaktır. Yapay zekâ taslağıdır; "
          "olaya katılan ekip gözden geçirmeden yayımlamayın."]
    return "\n".join(l) + "\n"


def tablo(ws, bas: list[str], satirlar: list[list], ai: set[int] = frozenset(), genis: dict | None = None) -> None:
    ws.append(bas)
    for h in ws[1]:
        h.fill, h.font, h.alignment = BASLIK_DOLGU, BASLIK_YAZI, UST
    for r in satirlar:
        ws.append(r)
        for j, v in enumerate(r, 1):
            ws.cell(ws.max_row, j).alignment = UST
            if j in ai and v not in (None, ""):
                ws.cell(ws.max_row, j).fill = AI_DOLGU
    for j in range(1, len(bas) + 1):
        ws.column_dimensions[get_column_letter(j)].width = (genis or {}).get(j, 16)
    ws.freeze_panes = "A2"


def rapor_yaz(cikti: Path, o: dict, c: dict, m: list, kalan: list[str], tur: int) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Zaman Çizelgesi"
    tablo(ws, ["Zaman", "Tür", "Olay"], [[datetime.strptime(x["zaman"].strip(), DAMGA) if x["zaman"].strip() in {h["damga"] for h in o["satirlar"]} else x["zaman"],
                                         x["tur"], x["olay"]] for x in c["zaman_cizelgesi"]], ai={2, 3}, genis={1: 18, 2: 18, 3: 90})
    for i in range(2, ws.max_row + 1):
        ws.cell(i, 1).number_format = "DD.MM.YYYY HH:MM"
    g = wb.create_sheet("Aksiyonlar")
    tablo(g, ["Olay", "Aksiyon", "Tür", "Sahip", "Termin", "Durum"],
          [[o["no"], x["aksiyon"], x["tur"], x["sahip"], datetime.strptime(x["termin"], "%d.%m.%Y") if re.fullmatch(r"\d{2}\.\d{2}\.\d{4}", x["termin"].strip()) else "",
            "Açık"] for x in c["aksiyonlar"]], ai={2, 3, 4}, genis={2: 70, 4: 20})
    for i in range(2, g.max_row + 1):
        g.cell(i, 5).number_format = "DD.MM.YYYY"
    g = wb.create_sheet("Süreler")
    tablo(g, ["Ölçü", "Değer", "Not"], [list(x) for x in m], genis={1: 36, 3: 60})
    for i in range(2, g.max_row + 1):
        if g.cell(i, 3).value == "Aşıldı":
            g.cell(i, 3).fill = KIRMIZI
    g = wb.create_sheet("Kontrol")
    g.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    g.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    g.cell(2, 1).fill = KIRMIZI if kalan else YESIL
    g.append(["Mor hücreler modelin metinleridir; süreler koddur. Rapor metni aynı adlı .md dosyasındadır. Olaya katılan ekip gözden "
              "geçirmeden yayımlamayın; kişi adı içermediğini kontrol edin."])
    g.column_dimensions["A"].width = 140
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    md = cikti.with_suffix(".md")
    md.write_text(rapor_metni(o, c, m), encoding="utf-8")
    return md


def calistir(yol: Path, cikti: Path, sla_dk: int | None = None, gizle: list[str] | None = None, tur: int = 2, evet: bool = False) -> dict:
    o = notlari_oku(yol, gizle)
    print(f"[OK] {o['no']} · {len(o['satirlar'])} zaman damgalı satır · {len(o['aksiyonlar'])} aksiyon · {len(o['kisiler'])} kişi maskelendi")
    llm.onay_al(f"Olay notları gönderilecek ({len(o['satirlar'])} satır; kişi adları takma adla, telefon, e-posta, TCKN, IBAN maskeli). "
                "Sunucu adı, IP ve sistem bilgileri gönderilir; gizliyse notlardan çıkarın.", evet)
    c, kalan, yapilan = duzeltme_dongusu(lambda d, onceki: uret(o, d, onceki), lambda x: denetle(o, x), tur)
    m = metrikler(c, sla_dk)
    md = rapor_yaz(cikti, o, c, m, kalan, yapilan)
    return {"olay": o, "rapor": c, "metrik": m, "kalan": kalan, "tur": yapilan, "md": md}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Olay notlarından suçlamasız BT olay sonrası raporu hazırlar; süreleri kod hesaplar.")
    p.add_argument("--notlar", type=Path, default=ORNEK / "olay_notlari.txt", help="Olay notları (.txt/.docx/.pdf): 'Olay: No | Hizmet: … | Tarih: GG.AA.YYYY' + 'SS:DD konuşan: mesaj'")
    p.add_argument("--sla-dk", type=int, help="Kesinti için SLA hedefi (dakika)")
    p.add_argument("--gizle", nargs="*", default=[], metavar="AD", help="Mesaj içinde geçen ve gönderilmeyecek ek adlar ('Ad Soyad:' konuşanlar zaten maskelenir)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "olay_sonrasi_raporu.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    sla = a.sla_dk if a.sla_dk is not None else (60 if a.notlar == ORNEK / "olay_notlari.txt" else None)
    try:
        s = calistir(a.notlar, a.cikti, sla, a.gizle, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    for ad, v, n in s["metrik"]:
        if v is not None:
            print(f"[OK] {ad}: {v}" + (f" · {n}" if n in ("Aşıldı", "Karşılandı") else ""))
    print(f"[OK] Kök neden ({s['rapor']['kok_neden_durumu']}) · {len(s['rapor']['aksiyonlar'])} aksiyon · {len(s['rapor']['acik_sorular'])} açık soru")
    for x in s["kalan"]:
        print(f"[!] Kalan: {x}")
    print(f"[OK] Rapor: {s['md'].resolve()} (+ {a.cikti.name})")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
