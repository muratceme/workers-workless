"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import json
import os
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri" / "olay_notlari.txt"
CEVAP = json.loads((Path(__file__).parent / "sahte_cevaplar.json").read_text(encoding="utf-8"))
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    """İlk turda katkıda bulunan etkenlerde bir kişiyi takma adıyla suçlar."""
    MESAJLAR.append(kullanici)
    return copy.deepcopy(CEVAP["dogru" if "<duzeltme>" in kullanici else "hatali"])


class KodTesti(unittest.TestCase):
    def test_okuma(self):
        o = agent.notlari_oku(ORNEK)
        self.assertEqual((o["no"], o["hizmet"], len(o["satirlar"]), len(o["aksiyonlar"])), ("INC-2026-0412", "ERP", 8, 3))
        self.assertEqual(o["kisiler"], {"Gülay Tekin": "[KİŞİ-1]", "Ozan Er": "[KİŞİ-2]", "Selim Ak": "[KİŞİ-3]"})
        self.assertIn("[14.10.2026 08:40] İzleme: erp-db01", o["metin"])               # sistem kaynağı maskelenmez
        for ad in o["kisiler"]:
            self.assertNotIn(ad, o["metin"])
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "n.txt"
            y.write_text("Olay: X | Tarih: 31.12.2026\n23:50 Nöbetçi: uyarı\n00:10 Nöbetçi: düzeldi\n", encoding="utf-8")
            self.assertEqual([h["damga"] for h in agent.notlari_oku(y)["satirlar"]], ["31.12.2026 23:50", "01.01.2027 00:10"])  # gece yarısı
            y.write_text("08:00 a: b\n", encoding="utf-8")
            with self.assertRaises(llm.LLMHatasi):
                agent.notlari_oku(y)

    def test_denetim(self):
        o = agent.notlari_oku(ORNEK)
        self.assertEqual(agent.denetle(o, CEVAP["dogru"]), [])
        self.assertEqual(agent.denetle(o, CEVAP["hatali"]), ["raporda kişi takma adı var ([KİŞİ-3]); suçlamasız yazın, rol / ekip kullanın"])
        c = copy.deepcopy(CEVAP["dogru"])
        z = c["zaman_cizelgesi"]
        z[4]["zaman"] = "14.10.2026 09:30"                                  # notlarda yok
        z[2]["tur"] = "Bilgi"                                               # tespit yok
        z[5], z[6] = z[6], z[5]                                             # sıra bozuk
        c["etki"] += " Kesinti 21 dakika sürdü."
        c["aksiyonlar"].pop(1)
        c["aksiyonlar"].append({"aksiyon": "Tüm sunuculara disk temizleme betiği kurulsun", "tur": "Önleyici", "sahip": "", "termin": "30.11.2026"})
        s = agent.denetle(o, c)
        for p in ("olmayan damga: 14.10.2026 09:30", "sıralı değil", "'Tespit' zaman çizelgesinde 0 kez", "olmayan sayı: 21", "olmayan tarih: 30.11.2026",
                  "rapora alınmamış: \"Servis hesabı parola", "notlarda olmayan aksiyon: \"Tüm sunuculara", "sahibi boş"):
            self.assertTrue(any(p in x for x in s), p)
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "n.txt"
            y.write_text(ORNEK.read_text(encoding="utf-8").replace(" Kök neden doğrulandı.", ""), encoding="utf-8")
            self.assertIn("durum 'Varsayım' olmalı", agent.denetle(agent.notlari_oku(y), CEVAP["dogru"])[0])

    def test_metrikler(self):
        m = {a: (v, n) for a, v, n in agent.metrikler(CEVAP["dogru"], 15)}
        self.assertEqual([m[k][0] for k in ("İlk belirtiden tespite (dk)", "Tespitten hizmetin düzelmesine (dk)", "Kesinti süresi (dk)", "Kalıcı çözüme kadar (dk)")],
                         [25, 21, 21, 85])
        self.assertEqual(m["SLA hedefi (dk)"], (15, "Aşıldı"))
        c = copy.deepcopy(CEVAP["dogru"])
        c["zaman_cizelgesi"] = [x for x in c["zaman_cizelgesi"] if x["tur"] != "Hizmet düzeldi"]
        m = {a: (v, n) for a, v, n in agent.metrikler(c, 60)}
        self.assertEqual((m["Kesinti süresi (dk)"][0], m["SLA hedefi (dk)"][1]), (None, "Kesinti süresi hesaplanamadı"))
        self.assertIn("Durum", m)


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ornek_ve_duzeltme(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "o.xlsx"
            self.assertEqual(agent.main(["--evet", "--cikti", str(cikti)]), 0)
            self.assertEqual(len(MESAJLAR), 2)
            for m in MESAJLAR:
                for ad in ("Gülay Tekin", "Ozan Er", "Selim Ak"):
                    self.assertNotIn(ad, m)
            wb = load_workbook(cikti)
            self.assertEqual(wb["Zaman Çizelgesi"]["A2"].value, datetime(2026, 10, 14, 8, 40))
            a = [[c.value for c in r] for r in wb["Aksiyonlar"].iter_rows(min_row=2)]
            self.assertEqual([(x[3], x[4]) for x in a], [("Altyapı ekibi", datetime(2026, 10, 31)), ("BT Güvenlik", None), ("Sistem ekibi", datetime(2026, 10, 21))])
            md = cikti.with_suffix(".md").read_text(encoding="utf-8")
            self.assertIn("- Kesinti süresi (dk): 21", md)
            self.assertIn("SLA hedefi (dk): 60 · Karşılandı", md)
            self.assertNotIn("[KİŞİ-", md)

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(agent.main(["--notlar", str(Path(tmp) / "yok.txt"), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
