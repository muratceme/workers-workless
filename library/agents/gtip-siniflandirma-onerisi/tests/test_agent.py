"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

V = KLASOR / "ornek_veri"
URUN, TAR = V / "urunler.csv", V / "tarife.csv"
MESAJLAR = []


def o(kod, gyk, oz, gerekce, guven):
    return {"kod": kod, "gyk": gyk, "kullanilan_ozellikler": oz, "gerekce": gerekce, "guven": guven}


def cevap(dogru=True) -> dict:
    t = {"urunler": [
        {"urun": "U-101", "oneriler": [o("3923.21", ["GYK 1", "GYK 6"], ["Alçak yoğunluklu polietilen (LDPE)", "Mağazalarda satılan ürünlerin taşınması"],
                                         "Eşya taşıma amaçlı torba; etilen polimerinden olduğu için 3923.21.", "Yüksek")], "eksik_bilgi": []},
        {"urun": "U-102", "oneriler": [o("3924.10", ["GYK 1", "GYK 6"], ["Evde mutfakta yiyecek saklama"], "Ev mutfağında kullanılan kap; 39. fasıl.", "Orta"),
                                       o("3923.10", ["GYK 1", "GYK 3(a)", "GYK 6"], ["perakende satış ambalajında"],
                                         "Ambalajlama amaçlı kutu olarak değerlendirilebilir, ancak kullanım ev tipi.", "Düşük")], "eksik_bilgi": []},
        {"urun": "U-103", "oneriler": [o("3923.30", ["GYK 1", "GYK 6"], ["500 ml, boş, kapaksız şişe"], "Dolum için boş şişe.", "Yüksek")], "eksik_bilgi": []},
        {"urun": "U-104", "oneriler": [o("4819.10", ["GYK 1", "GYK 6"], ["3 katlı oluklu mukavvadan"], "Oluklu kartondan koli.", "Orta")],
         "eksik_bilgi": ["Malzemesi kağıt esaslı oluklu karton mu?"]},
        {"urun": "U-105", "oneriler": [o("3923.50", ["GYK 1", "GYK 6"], ["28 mm vidalı şişe kapağı"], "Plastikten kapama malzemesi.", "Orta")],
         "eksik_bilgi": ["Kapak hangi şişelerde ve hangi amaçla kullanılıyor?"]}]}
    if not dogru:
        t["urunler"][0]["oneriler"][0]["kod"] = "3923.21.00.00.11"
        t["urunler"][2]["oneriler"] = [o("3923", ["GYK 1"], ["şeffaf"], "Plastik şişe, 1,5 litre.", "Yüksek")]
        t["urunler"][3]["oneriler"][0]["guven"] = "Yüksek"
        t["urunler"][3]["eksik_bilgi"] = []
        del t["urunler"][4]
    return t


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return cevap("<duzeltme>" in kullanici)


class KodTesti(unittest.TestCase):
    def test_okuma(self):
        u, t = agent.urunleri_oku(URUN), agent.tarife_oku(TAR)
        self.assertEqual(([x["kod"] for x in u], u[3]["eksik"], u[4]["eksik"]), (["U-101", "U-102", "U-103", "U-104", "U-105"], ["Malzeme"], ["Kullanım Amacı"]))
        self.assertEqual((len(t), t["3923"]["tanim"][:10]), (15, "Eşya taşım"))
        self.assertEqual(agent.ust_pozisyon("3923.21", t)["kod"], "3923")

    def test_denetim(self):
        u, t = agent.urunleri_oku(URUN), agent.tarife_oku(TAR)
        self.assertEqual(agent.denetle(cevap(), u, t), [])
        s = agent.denetle(cevap(False), u, t)
        for parca in ("U-101 / 3923.21.00.00.11: kod tarife tablosunda yok", "alt pozisyonları var", "'şeffaf'", "1.5", "U-104: Malzeme eksik; eksik bilgi",
                      "U-104: Malzeme eksikken", "U-105: ürün için öneri yok"):
            self.assertTrue(any(parca in x for x in s), (parca, s))
        self.assertEqual(len(s), 7, s)


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "g.xlsx"
            s = agent.calistir(URUN, TAR, cikti, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            self.assertIn("3923.30 | Damacanalar", MESAJLAR[0])
            ws = load_workbook(cikti)["Öneriler"]
            self.assertEqual((ws["D3"].value, ws["E3"].value, ws["F3"].value[:4]), ("3924.10", "Sofra ve mutfak eşyası", "3924"))
            self.assertIn("Bağlayıcı Tarife Bilgisi", ws.cell(ws.max_row, 1).value)

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(URUN, TAR, Path(tmp) / "g.xlsx", tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 7)
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--evet", "--tarife", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
