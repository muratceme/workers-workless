"""
GTİP Sınıflandırma Önerisi — Gerisi Bende AI Agent
Dış Ticaret › Dış Ticaret Uzmanı

1. Kod ürün listesini (ürün kodu, adı, tanımı, malzemesi, kullanım amacı) ve kullanıcının verdiği tarife tablosunu
   (Türk Gümrük Tarife Cetveli'nden ilgili fasıllar: kod; tanım) okur.
2. Model her ürün için tarife tablosundaki kodlardan en fazla 3 öneriyi; uyguladığı Genel Yorum Kurallarını (GYK),
   dayandığı ürün özelliklerini, gerekçesini ve güven düzeyini yazar; sınıflandırma için eksik bilgileri sorar.
3. Kod denetler:
   - Önerilen kod tarife tablosunda var mı (kod uydurulamaz); alt pozisyonu olan bir pozisyon seçilmişse alt pozisyon istenir
   - GYK 1 uygulanmış mı; 4 haneden uzun kodlarda GYK 6 belirtilmiş mi
   - Gerekçede kullanılan özellikler ürün bilgisinde aynen geçiyor mu; gerekçedeki sayılar ürün bilgisinde ya da tabloda var mı
   - Malzeme ya da kullanım amacı eksik ürünlerde eksik bilgi sorulmuş ve "Yüksek" güven verilmemiş mi; birden fazla
     "Yüksek" güvenli öneri var mı
   Sorun varsa modele geri gönderilir (en fazla --tur kez).
4. Kod her öneriye üst pozisyonun tanımını ekler ve Excel raporu yazar. Sonuç öneridir; gümrük müşaviri onayı gerekir.

Kullanım:
    python agent.py                                                  # örnek (kurgusal) ambalaj ürünleri
    python agent.py --urunler urunler.xlsx --tarife fasil39_48.xlsx
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
URUN = {"kod": ("Ürün Kodu", "Kod", "Stok Kodu"), "ad": ("Ürün Adı", "Ad"), "tanim": ("Tanım", "Açıklama", "Ürün Tanımı"),
        "malzeme": ("Malzeme", "Hammadde"), "kullanim": ("Kullanım Amacı", "Kullanım")}
TARIFE = {"kod": ("Kod", "GTİP", "Pozisyon"), "tanim": ("Tanım", "Eşya Tanımı", "Açıklama")}
GYK = ["GYK 1", "GYK 2(a)", "GYK 2(b)", "GYK 3(a)", "GYK 3(b)", "GYK 3(c)", "GYK 4", "GYK 5(a)", "GYK 5(b)", "GYK 6"]
GUVEN = ["Yüksek", "Orta", "Düşük"]

SEMA = {
    "type": "object",
    "properties": {
        "urunler": {"type": "array", "items": {"type": "object", "properties": {
            "urun": S,
            "oneriler": {"type": "array", "items": {"type": "object", "properties": {
                "kod": S, "gyk": {"type": "array", "items": {"type": "string", "enum": GYK}}, "kullanilan_ozellikler": {"type": "array", "items": S},
                "gerekce": S, "guven": {"type": "string", "enum": GUVEN}},
                "required": ["kod", "gyk", "kullanilan_ozellikler", "gerekce", "guven"], "additionalProperties": False}},
            "eksik_bilgi": {"type": "array", "items": S}},
            "required": ["urun", "oneriler", "eksik_bilgi"], "additionalProperties": False}},
    },
    "required": ["urunler"],
    "additionalProperties": False,
}


def rakam(kod) -> str:
    return re.sub(r"\D", "", str(kod or ""))


def tarife_oku(yol: Path) -> dict:
    t = {}
    for r in kayitlar(yol, TARIFE, ("kod", "tanim")):
        if rakam(r["kod"]):
            t[rakam(r["kod"])] = {"kod": str(r["kod"]).strip(), "tanim": str(r["tanim"]).strip()}
    if not t:
        raise llm.LLMHatasi(f"{yol.name}: tarife tablosu boş.")
    return t


def urunleri_oku(yol: Path) -> list[dict]:
    out = []
    for r in kayitlar(yol, URUN, ("ad", "tanim")):
        if not (r["ad"] or r["tanim"]):
            continue
        r = {k: str(v) for k, v in r.items()}
        r["kod"] = r["kod"] or f"U{len(out) + 1}"
        r["eksik"] = [URUN[k][0] for k in ("malzeme", "kullanim") if not r[k]]
        out.append(r)
    if not out:
        raise llm.LLMHatasi(f"{yol.name}: ürün yok.")
    return out


def ust_pozisyon(kod: str, tarife: dict) -> dict | None:
    r = rakam(kod)
    for n in range(len(r) - 1, 3, -1):
        if r[:n] in tarife and r[:n] != r:
            return tarife[r[:n]]
    return tarife.get(r[:4]) if r[:4] != r else None


def urun_metni(u: dict) -> str:
    return " ".join(u[k] for k in ("ad", "tanim", "malzeme", "kullanim"))


def denetle(t: dict, urunler: list[dict], tarife: dict) -> list[str]:
    sorun = []
    u = {x["kod"]: x for x in urunler}
    tablo = " ".join(f"{v['kod']} {v['tanim']}" for v in tarife.values())
    fasillar = {str(int(k[:2])) for k in tarife} | {k[:2] for k in tarife} | {k[:4] for k in tarife}
    gorulen = set()
    for c in t["urunler"]:
        if c["urun"] not in u:
            sorun.append(f"{c['urun']}: ürün listesinde yok")
            continue
        gorulen.add(c["urun"])
        x = u[c["urun"]]
        if not 1 <= len(c["oneriler"]) <= 3:
            sorun.append(f"{c['urun']}: 1 ile 3 arasında öneri olmalı ({len(c['oneriler'])})")
        yuksek = [o["kod"] for o in c["oneriler"] if o["guven"] == "Yüksek"]
        if len(yuksek) > 1:
            sorun.append(f"{c['urun']}: birden fazla 'Yüksek' güvenli öneri ({', '.join(yuksek)}); en olası tek kodu seçin")
        if x["eksik"]:
            if not c["eksik_bilgi"]:
                sorun.append(f"{c['urun']}: {', '.join(x['eksik'])} eksik; eksik bilgi sorulmalı")
            if yuksek:
                sorun.append(f"{c['urun']}: {', '.join(x['eksik'])} eksikken 'Yüksek' güven verilemez")
        for o in c["oneriler"]:
            ad = f"{c['urun']} / {o['kod']}"
            r = rakam(o["kod"])
            if r not in tarife:
                sorun.append(f"{ad}: kod tarife tablosunda yok (kod uydurulamaz; yalnız tablodaki kodları önerin)")
            elif any(k != r and k.startswith(r) for k in tarife):
                sorun.append(f"{ad}: bu pozisyonun alt pozisyonları var; alt pozisyon seçin (GYK 6)")
            if "GYK 1" not in o["gyk"]:
                sorun.append(f"{ad}: GYK 1 (pozisyon metinleri ve notlar) belirtilmemiş")
            if len(r) > 4 and "GYK 6" not in o["gyk"]:
                sorun.append(f"{ad}: alt pozisyon seçimi için GYK 6 belirtilmemiş")
            if not o["kullanilan_ozellikler"]:
                sorun.append(f"{ad}: dayanılan ürün özelliği yazılmamış")
            for oz in o["kullanilan_ozellikler"]:
                if not metinde_var(oz, urun_metni(x)):
                    sorun.append(f"{ad}: '{oz}' ürün bilgisinde geçmiyor")
            fazla = sorted(sayilar(o["gerekce"]) - sayilar(urun_metni(x) + " " + tablo) - fasillar - {"1", "2", "3", "4", "5", "6"})
            if fazla:
                sorun.append(f"{ad}: gerekçede ürün bilgisinde / tabloda olmayan sayı: {', '.join(fazla)}")
    for k in u:
        if k not in gorulen:
            sorun.append(f"{k}: ürün için öneri yok")
    return sorun


def uret(urunler: list[dict], tarife: dict, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    ur = "\n".join(f"{x['kod']} | {x['ad']} | {x['tanim']} | malzeme: {x['malzeme'] or '(belirtilmemiş)'} | kullanım: {x['kullanim'] or '(belirtilmemiş)'}"
                   for x in urunler)
    ta = "\n".join(f"{v['kod']} | {v['tanim']}" for v in tarife.values())
    mesaj = f"<urunler>\n{ur}\n</urunler>\n<tarife>\n{ta}\n</tarife>" + duzeltme_blogu(duzeltme, onceki)
    return llm.json_iste(sistem, mesaj, SEMA)


def rapor_yaz(cikti: Path, t, urunler, tarife, kalan, tur) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Öneriler"
    baslik(ws, ["Ürün Kodu", "Ürün Adı", "Sıra", "Önerilen Kod", "Tarife Tanımı", "Üst Pozisyon", "GYK", "Dayanılan Özellikler", "Gerekçe", "Güven",
                "Eksik Bilgi"], (12, 24, 6, 14, 40, 40, 18, 36, 60, 9, 40))
    u = {x["kod"]: x for x in urunler}
    for c in t["urunler"]:
        for i, o in enumerate(c["oneriler"], 1):
            ust = ust_pozisyon(o["kod"], tarife)
            ws.append([c["urun"], u.get(c["urun"], {}).get("ad", "?"), i, o["kod"], tarife.get(rakam(o["kod"]), {}).get("tanim", "TABLODA YOK"),
                       f"{ust['kod']} {ust['tanim']}" if ust else "", ", ".join(o["gyk"]), "\n".join(o["kullanilan_ozellikler"]), o["gerekce"], o["guven"],
                       "\n".join(c["eksik_bilgi"]) if i == 1 else ""])
            boya(ws, (9,), AI_DOLGU)
            boya(ws, (10,), {"Yüksek": YESIL, "Orta": SARI, "Düşük": KIRMIZI}[o["guven"]])
    son_not(ws, "Önemli: Öneridir, bağlayıcı değildir. Kesin sınıflandırma için gümrük müşaviri onayı alın; gerekirse Bağlayıcı Tarife Bilgisi (BTB) "
                "başvurusu yapın. Tarife tablosunun güncel olduğunu kontrol edin.", kalan, tur)
    hizala(ws)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(urun_yol: Path, tarife_yol: Path, cikti: Path, tur: int = 2, evet: bool = False) -> dict:
    urunler = urunleri_oku(urun_yol)
    tarife = tarife_oku(tarife_yol)
    print(f"[OK] {len(urunler)} ürün · tarife tablosunda {len(tarife)} kod")
    if len(tarife) > 2000:
        print("[!] Tarife tablosu çok büyük; yalnız ilgili fasılları verin (maliyet ve doğruluk için).")
    for x in urunler:
        if x["eksik"]:
            print(f"[!] {x['kod']}: {', '.join(x['eksik'])} eksik")
    llm.onay_al(f"{len(urunler)} ürün tanımı ve tarife tablosu gönderilecek.", evet)
    t, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(urunler, tarife, d, o), lambda x: denetle(x, urunler, tarife), tur)
    rapor_yaz(cikti, t, urunler, tarife, kalan, yapilan)
    return {"urunler": urunler, "taslak": t, "kalan": kalan, "tur": yapilan}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    p = argparse.ArgumentParser(description="Ürün tanımından, verdiğiniz tarife tablosundaki olası GTİP kodlarını gerekçesiyle önerir.")
    v = BURASI / "ornek_veri"
    p.add_argument("--urunler", type=Path, default=v / "urunler.csv", help="Ürün listesi (CSV/Excel): Ürün Kodu, Ürün Adı, Tanım, Malzeme, Kullanım Amacı")
    p.add_argument("--tarife", type=Path, default=v / "tarife.csv", help="Tarife tablosu (CSV/Excel): Kod, Tanım — Türk Gümrük Tarife Cetveli'nden ilgili fasıllar")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "gtip_onerileri.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.urunler, a.tarife, a.cikti, a.tur, a.evet)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    for c in s["taslak"]["urunler"]:
        if c["oneriler"]:
            print(f"[OK] {c['urun']}: {c['oneriler'][0]['kod']} ({c['oneriler'][0]['guven']})")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
