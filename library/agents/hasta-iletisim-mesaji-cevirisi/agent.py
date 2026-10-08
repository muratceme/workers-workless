"""
Hasta İletişim Mesajı Çevirisi — Workers / Workless AI Agent
Sağlık (Özel Hastane) › Uluslararası Hasta Hizmetleri › Uluslararası Hasta Koordinatörü

1. mesajlar/ klasöründeki hasta mesajları (.txt/.docx/.pdf/.eml; isteğe bağlı 'Hasta:' ve 'Tarih:' başlık
   satırları), kurum_bilgileri.txt ve terimce.csv okunur.
2. Hasta adları [HASTA-n] takma adlarıyla; telefon (uluslararası dahil), e-posta, pasaport numarası, TCKN ve IBAN
   maskelenir. Sağlık verisi özel nitelikli kişisel veri olduğundan gönderim öncesi açık onay istenir.
3. Model her mesaj için Türkçe çeviri, terimler, talepler, aciliyet (Acil / Öncelikli / Rutin), hastanın dilinde
   cevap taslağı, cevabın Türkçe karşılığı, koordinatörün dolduracağı yer tutucular ve doktora iletilecek tıbbi
   hususları yazar. Tıbbi tavsiye vermez; bilinmeyen fiyat, tarih ve süreleri uydurmaz.
4. Kod denetler:
   - Doz ve ölçüm değerleri (500 mg, 38,9 °C, 5 мг…) çeviride aynı sayı ve birimle geçiyor mu?
   - Terimcedeki kaynak terim mesajda geçiyorsa Türkçe karşılığı çeviride var mı?
   - Çeviride kırmızı bayrak ifadesi (ateş, nefes darlığı, akıntı, kanama…) varken aciliyet "Acil" değil mi?
     Acil mesajın cevabı acil servise yönlendiriyor mu?
   - Cevapta mesajda ve kurum bilgilerinde olmayan sayı (uydurma fiyat / tarih) var mı; cevap ile Türkçe karşılığı
     aynı sayıları içeriyor mu; doldurulmamış yer tutucular.
   - Mesajın alfabesi (Kiril, Arap) ile modelin bildirdiği dil uyumlu mu?
5. Çıktı: mesaj başına Markdown (takma adlar geri açılmış) + Excel (Mesajlar aciliyete göre, Terimler, Kontroller).

Kullanım:
    python agent.py                                          # örnek: 4 kurgusal hasta mesajı (EN, DE, RU)
    python agent.py --girdi ./gelen_kutusu --gizle "Refakatçi Adı"
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from dataclasses import dataclass, field
from email import policy
from email.parser import BytesParser
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
KIRMIZI_BAYRAK = ("ates", "nefes darligi", "nefes almakta zorlan", "gogus agrisi", "bilinc", "bayil", "kanama", "akinti", "intihar", "felc",
                  "siddetli agri", "nobet", "havale", "kusma", "baygin")
ALFABE = {"Kiril": (re.compile(r"[Ѐ-ӿ]"), {"ru", "uk", "bg", "sr", "kk", "be", "mk", "ky", "tg", "mn"}),
          "Arap": (re.compile(r"[؀-ۿ]"), {"ar", "fa", "ur", "ps", "ku"})}
BIRIM = {"mg": "mg", "мг": "mg", "mcg": "mcg", "µg": "mcg", "μg": "mcg", "мкг": "mcg", "g": "g", "г": "g", "ml": "ml", "мл": "ml", "l": "l", "л": "l",
         "iu": "U", "ie": "U", "me": "U", "ед": "U", "u": "U", "ü": "U", "ünite": "U", "units": "U", "unit": "U", "einheiten": "U",
         "°c": "°C", "°f": "°F", "mmhg": "mmHg", "ммрт": "mmHg", "kg": "kg", "кг": "kg", "cm": "cm", "см": "cm", "mm": "mm", "мм": "mm",
         "bpm": "bpm", "%": "%"}
OLCU_RE = re.compile(r"(\d+(?:[.,]\d+)?)\s*(мм рт|mmhg|einheiten|ünite|units?|mcg|мкг|µg|μg|°\s?[cf]|bpm|mg|мг|ml|мл|kg|кг|cm|см|mm|мм|iu|ie|me|ед|g|г|l|л|u|ü|%)"
                     r"(?![A-Za-zА-Яа-яçğıöşüÇĞİÖŞÜ])", re.I)
YER_TUTUCU = re.compile(r"\[[A-ZÇĞİÖŞÜ][A-ZÇĞİÖŞÜ0-9 _/.-]{1,40}\]")
MASKE_ADLARI = re.compile(r"^\[(HASTA|KİŞİ)-\d+\]$|^\[(TELEFON|E-POSTA|PASAPORT|TCKN|IBAN|LINKEDIN)\]$")
TELEFON_ULUSLARARASI = re.compile(r"\+\d{1,3}[\s.-]?\(?\d{1,4}\)?(?:[\s.-]?\d{2,4}){2,4}")
PASAPORT = re.compile(r"(?i)(passport|reisepass|pasaport|паспорт[а-я]*)(\s*(?:no\.?|number|nummer|nr\.?|номер|numarası)?\s*[:#]?\s*)"
                      r"([A-ZА-Я0-9][A-ZА-Я0-9 ]{4,12}\d)")


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def katla(s) -> str:
    s = kucuk(s)
    for a, b in zip("ıöüşğçâîû", "iousgcaiu"):
        s = s.replace(a, b)
    return re.sub(r"[^a-z0-9а-яё%]+", " ", s).strip()


def sayi_norm(s: str) -> str:
    s = s.replace(",", ".")
    return s.rstrip("0").rstrip(".") if "." in s else s


def olculer(metin: str) -> list[tuple[str, str]]:
    sonuc = []
    for n, b in OLCU_RE.findall(metin):
        b = b.lower().replace(" ", "")
        sonuc.append((sayi_norm(n), BIRIM.get(b, b)))
    return sonuc


def rakamlar(metin: str) -> set[str]:
    metin = YER_TUTUCU.sub(" ", metin)
    return {sayi_norm(x) for x in re.findall(r"\d+(?:[.,]\d+)?", metin)}


# ----------------------------------------------------------------------------
# Girdiler
# ----------------------------------------------------------------------------

@dataclass
class Mesaj:
    no: str
    dosya: str
    hasta: list[str]
    tarih: str
    metin: str
    yanit: dict = field(default_factory=dict)
    sorunlar: list = field(default_factory=list)


def eml_oku(yol: Path) -> tuple[str, str, str]:
    m = BytesParser(policy=policy.default).parse(yol.open("rb"))
    govde = m.get_body(preferencelist=("plain", "html"))
    metin = govde.get_content() if govde else ""
    if govde is not None and govde.get_content_type() == "text/html":
        metin = re.sub(r"<[^>]+>", " ", metin)
    gonderen = re.sub(r"<[^>]*>", "", str(m.get("From", ""))).strip().strip('"')
    return metin, gonderen, str(m.get("Date", ""))


def mesaj_oku(yol: Path) -> Mesaj:
    if yol.suffix.lower() == ".eml":
        metin, gonderen, tarih = eml_oku(yol)
        return Mesaj(yol.stem, yol.name, [gonderen] if gonderen else [], tarih, metin.strip())
    metin = belge.metin_oku(yol)
    hasta, tarih, govde = [], "", []
    baslik = True
    for satir in metin.splitlines():
        if baslik:
            m = re.match(r"^\s*(Hasta|Patient|Ad Soyad|Tarih|Date)\s*:\s*(.*)$", satir, re.I)
            if m:
                if katla(m.group(1)) in ("hasta", "patient", "ad soyad"):
                    hasta += [x.strip() for x in m.group(2).split(";") if x.strip()]
                else:
                    tarih = m.group(2).strip()
                continue
            if not satir.strip() and not govde:
                continue
            baslik = False
        govde.append(satir)
    return Mesaj(yol.stem, yol.name, hasta, tarih, "\n".join(govde).strip())


def terimce_oku(yol: Path | None) -> list[tuple[str, str]]:
    if yol is None or not yol.exists():
        return []
    metin = belge.txt_oku(yol)
    ilk = "\n".join(metin.splitlines()[:5])
    satirlar = [r for r in csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)) if len(r) >= 2 and r[0].strip()]
    if satirlar and katla(satirlar[0][0]) in ("kaynak terim", "kaynak", "terim"):
        satirlar = satirlar[1:]
    return [(r[0].strip(), r[1].strip()) for r in satirlar if r[1].strip()]


def dosyalari_oku(klasor: Path) -> tuple[list[Mesaj], str, list[tuple[str, str]], list[str]]:
    mesajlar, atlanan = [], []
    kaynak = klasor / "mesajlar" if (klasor / "mesajlar").is_dir() else klasor
    for yol in sorted(p for p in kaynak.iterdir() if p.is_file()):
        ad = katla(yol.stem)
        if kaynak == klasor and (ad.startswith(("kurum", "terimce")) or yol.suffix.lower() == ".csv"):
            continue
        if yol.suffix.lower() not in belge.DESTEKLENEN | {".eml"}:
            atlanan.append(yol.name)
            continue
        try:
            m = mesaj_oku(yol)
        except (belge.BelgeHatasi, OSError, ValueError) as h:
            atlanan.append(f"{yol.name} ({h})")
            continue
        if m.metin:
            mesajlar.append(m)
    kurum = ""
    for yol in sorted(klasor.glob("kurum*")):
        try:
            kurum += belge.metin_oku(yol) + "\n"
        except belge.BelgeHatasi as h:
            atlanan.append(f"{yol.name} ({h})")
    terimce = terimce_oku(next(iter(sorted(klasor.glob("terimce*.csv"))), None))
    return mesajlar, kurum.strip(), terimce, atlanan


# ----------------------------------------------------------------------------
# Maskeleme
# ----------------------------------------------------------------------------

def takma_adlar(mesajlar: list[Mesaj], terimler: list[str]) -> dict[str, str]:
    harita = {}
    n = 0
    for m in mesajlar:
        if not m.hasta:
            continue
        mevcut = next((harita[a] for a in m.hasta if a in harita), None)
        if mevcut is None:
            n += 1
            mevcut = f"[HASTA-{n}]"
        for ad in m.hasta:
            harita.setdefault(ad, mevcut)
            for parca in ad.split():
                if len(parca) >= 3:
                    harita.setdefault(parca, mevcut)
    for i, ad in enumerate(dict.fromkeys(t.strip() for t in terimler if t.strip()), 1):
        harita.setdefault(ad, f"[KİŞİ-{i}]")
    return harita


def maskele(metin: str, harita: dict[str, str]) -> str:
    metin = PASAPORT.sub(lambda m: m.group(1) + m.group(2) + "[PASAPORT]", metin)
    metin = TELEFON_ULUSLARARASI.sub("[TELEFON]", metin)
    for gercek, takma in sorted(harita.items(), key=lambda x: -len(x[0])):
        metin = re.sub(r"(?<!\w)" + re.escape(gercek) + r"(?!\w)", takma, metin, flags=re.I)
    return llm.maskele(metin)


def geri_ac(metin: str, harita: dict[str, str]) -> str:
    ters = {}
    for gercek, takma in harita.items():          # ilk eklenen tam addır
        ters.setdefault(takma, gercek)
    for takma, gercek in sorted(ters.items(), key=lambda x: -len(x[0])):
        metin = metin.replace(takma, gercek)
    return metin


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

SEMA = {
    "type": "object",
    "properties": {
        "dil": {"type": "string"},
        "dil_adi": {"type": "string"},
        "ceviri_tr": {"type": "string"},
        "terimler": {"type": "array", "items": {"type": "object", "properties": {"kaynak": {"type": "string"}, "turkce": {"type": "string"}},
                                                 "required": ["kaynak", "turkce"], "additionalProperties": False}},
        "talepler": {"type": "array", "items": {"type": "string"}},
        "aciliyet": {"type": "string", "enum": ["Acil", "Öncelikli", "Rutin"]},
        "aciliyet_gerekcesi": {"type": "string"},
        "cevap": {"type": "string"},
        "cevap_tr": {"type": "string"},
        "koordinatore_sorular": {"type": "array", "items": {"type": "string"}},
        "doktora_iletilecek": {"type": "string"},
    },
    "required": ["dil", "dil_adi", "ceviri_tr", "terimler", "talepler", "aciliyet", "aciliyet_gerekcesi", "cevap", "cevap_tr", "koordinatore_sorular",
                 "doktora_iletilecek"],
    "additionalProperties": False,
}
SIRA = {"Acil": 0, "Öncelikli": 1, "Rutin": 2}


def hazirla(m: Mesaj, kurum: str, terimce, harita) -> tuple[str, str]:
    mesaj = "\n".join([
        "<mesaj>", maskele(m.metin, harita), "</mesaj>",
        "<kurum_bilgileri>", kurum or "Verilmedi; kurum bilgisi gerektiren her şey için yer tutucu kullan.", "</kurum_bilgileri>",
        "<terimce>", *([f"{k} = {t}" for k, t in terimce] or ["Verilmedi."]), "</terimce>",
    ])
    return (BURASI / "prompt.md").read_text(encoding="utf-8"), mesaj


def denetle(m: Mesaj, y: dict, kurum: str, terimce, gonderilen: str) -> list[tuple[str, str]]:
    sorunlar = []
    # Ölçü / doz koruma
    ceviri_olcu = olculer(y["ceviri_tr"])
    for o in olculer(gonderilen):
        if o in ceviri_olcu:
            ceviri_olcu.remove(o)
        else:
            sorunlar.append(("yüksek", f"Doz / ölçü değeri çeviride aynı sayı ve birimle yok: {o[0]} {o[1]}"))
    # Terimce
    kaynak_k = kucuk(gonderilen)
    ceviri_k = katla(y["ceviri_tr"])
    for kaynak, turkce in terimce:
        if kucuk(kaynak) in kaynak_k and katla(turkce) not in ceviri_k:
            sorunlar.append(("orta", f"Terimce: '{kaynak}' mesajda geçiyor, çeviride '{turkce}' karşılığı yok"))
    # Alfabe ve dil
    for ad, (desen, diller) in ALFABE.items():
        if len(desen.findall(gonderilen)) > 10 and y["dil"].lower()[:2] not in diller:
            sorunlar.append(("orta", f"Mesaj {ad} alfabesiyle yazılmış ama model dili '{y['dil']}' bildirdi"))
    # Aciliyet
    bayrak = [k for k in KIRMIZI_BAYRAK if re.search(r"\b" + k, ceviri_k)]
    if bayrak and y["aciliyet"] != "Acil":
        sorunlar.append(("yüksek", f"Çeviride kırmızı bayrak ifadesi var ({', '.join(bayrak)}) ama aciliyet '{y['aciliyet']}'; doktor / koordinatör "
                                   "hemen değerlendirmeli"))
    if y["aciliyet"] == "Acil" and "acil" not in katla(y["cevap_tr"]):
        sorunlar.append(("yüksek", "Acil mesajın cevabı hastayı acil servise / acil numaraya yönlendirmiyor"))
    # Cevaptaki sayılar
    kaynak_sayi = rakamlar(gonderilen) | rakamlar(kurum)
    uydurma = sorted(rakamlar(y["cevap"]) - kaynak_sayi, key=lambda s: float(s))
    if uydurma:
        sorunlar.append(("yüksek", f"Cevapta mesajda ve kurum bilgilerinde olmayan sayı: {', '.join(uydurma)} (uydurma fiyat / tarih / süre olabilir)"))
    if rakamlar(y["cevap"]) != rakamlar(y["cevap_tr"]):
        fark = sorted(rakamlar(y["cevap"]) ^ rakamlar(y["cevap_tr"]))
        sorunlar.append(("orta", f"Cevap ile Türkçe karşılığındaki sayılar farklı: {', '.join(fark)}"))
    y["yer_tutucular"] = sorted({t for t in YER_TUTUCU.findall(y["cevap"]) if not MASKE_ADLARI.match(t)})
    if y["yer_tutucular"]:
        sorunlar.append(("bilgi", f"Gönderimden önce doldurulacak: {', '.join(y['yer_tutucular'])}"))
    if y["doktora_iletilecek"].strip():
        sorunlar.append(("bilgi", "Doktora iletilecek tıbbi husus var"))
    return sorunlar


def model_yaz(mesajlar, kurum, terimce, harita) -> None:
    for m in mesajlar:
        sistem, gonderilen = hazirla(m, kurum, terimce, harita)
        y = llm.json_iste(sistem, gonderilen, SEMA)
        govde = gonderilen.split("<mesaj>\n", 1)[1].split("\n</mesaj>", 1)[0]
        m.yanit = y
        m.sorunlar = denetle(m, y, kurum, terimce, govde)
        print(f"[OK] {m.no}: {y['dil_adi']} · {y['aciliyet']}")


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
ONEM = {"yüksek": "FDE2E1", "orta": "FFF4CE", "bilgi": "E8F0FE", "Acil": "FDE2E1", "Öncelikli": "FFF4CE", "Rutin": "E8F0FE"}
KARAR = PatternFill("solid", fgColor="FFF4CE")
MODEL = PatternFill("solid", fgColor="EDE7FF")
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def rapor_yaz(cikti: Path, mesajlar: list[Mesaj], harita, atlanan) -> Path:
    g = lambda s: geri_ac(s, harita)  # noqa: E731
    sirali = sorted(mesajlar, key=lambda m: (SIRA[m.yanit["aciliyet"]], m.no))
    wb = Workbook()
    ms = wb.active
    ms.title = "Mesajlar"
    _baslik(ms, ["Mesaj", "Hasta", "Tarih", "Dil", "Aciliyet", "Aciliyet Gerekçesi", "Türkçe Çeviri", "Talepler", "Cevap Taslağı", "Cevabın Türkçesi",
                 "Doldurulacak", "Doktora İletilecek", "Kontrol", "Koordinatör Onayı"], (8, 16, 15, 10, 10, 30, 60, 34, 60, 60, 20, 34, 30, 16))
    for m in sirali:
        y = m.yanit
        ms.append([m.no, ", ".join(m.hasta[:1]), m.tarih, y["dil_adi"], y["aciliyet"], g(y["aciliyet_gerekcesi"]), g(y["ceviri_tr"]),
                   "\n".join(f"- {g(t)}" for t in y["talepler"]), g(y["cevap"]), g(y["cevap_tr"]), ", ".join(y["yer_tutucular"]) or "—",
                   g(y["doktora_iletilecek"]) or "—", "\n".join(f"[{o}] {a}" for o, a in m.sorunlar if o != "bilgi") or "—", ""])
        r = ms.max_row
        ms.cell(r, 5).fill = PatternFill("solid", fgColor=ONEM[y["aciliyet"]])
        for j in (7, 9, 10):
            ms.cell(r, j).fill = MODEL
        if any(o == "yüksek" for o, _ in m.sorunlar):
            ms.cell(r, 13).fill = PatternFill("solid", fgColor=ONEM["yüksek"])
        ms.cell(r, 14).fill = KARAR
        for h in ms[r]:
            h.alignment = UST

    tr = wb.create_sheet("Terimler")
    _baslik(tr, ["Mesaj", "Kaynak", "Türkçe"], (8, 34, 34))
    for m in mesajlar:
        for t in m.yanit["terimler"]:
            tr.append([m.no, t["kaynak"], t["turkce"]])

    kt = wb.create_sheet("Kontroller")
    _baslik(kt, ["Mesaj", "Önem", "Kontrol", "İnceleme"], (8, 9, 100, 24))
    for m in sirali:
        for o, a in m.sorunlar:
            kt.append([m.no, o, g(a), ""])
            kt.cell(kt.max_row, 2).fill = PatternFill("solid", fgColor=ONEM[o])
            kt.cell(kt.max_row, 4).fill = KARAR
    kt.append([])
    kt.append(["", "", f"Model: {llm.kullanim_ozeti()} · okunamayan: {'; '.join(atlanan) or '—'}"])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

    klasor = cikti.parent / (cikti.stem + "_cevaplar")
    klasor.mkdir(parents=True, exist_ok=True)
    for m in sirali:
        y = m.yanit
        sat = [f"# {m.no} · {', '.join(m.hasta[:1]) or 'hasta'} · {y['dil_adi']} · **{y['aciliyet']}**", "",
               "> Taslaktır. Koordinatör (tıbbi içerikte doktor) kontrol etmeden ve yer tutucular doldurulmadan gönderilmez.", "",
               f"**Aciliyet gerekçesi:** {g(y['aciliyet_gerekcesi'])}", "", "## Türkçe çeviri", "", g(y["ceviri_tr"]), "", "## Talepler", ""]
        sat += [f"- {g(t)}" for t in y["talepler"]] + ["", f"## Cevap taslağı ({y['dil_adi']})", "", g(y["cevap"]), "", "## Cevabın Türkçesi", "",
                                                     g(y["cevap_tr"]), ""]
        if y["koordinatore_sorular"]:
            sat += ["## Koordinatörün dolduracakları", ""] + [f"- {g(s)}" for s in y["koordinatore_sorular"]] + [""]
        if y["doktora_iletilecek"].strip():
            sat += ["## Doktora iletilecek", "", g(y["doktora_iletilecek"]), ""]
        if m.sorunlar:
            sat += ["## Kontroller", ""] + [f"- **{o}** {g(a)}" for o, a in m.sorunlar] + [""]
        (klasor / f"{m.no}.md").write_text("\n".join(sat) + "\n", encoding="utf-8")
    return klasor


def calistir(klasor: Path, cikti: Path, terimler: list[str] | None = None, evet: bool = False) -> dict:
    if not klasor.is_dir():
        raise llm.LLMHatasi(f"{klasor} klasörü bulunamadı.")
    mesajlar, kurum, terimce, atlanan = dosyalari_oku(klasor)
    if not mesajlar:
        raise llm.LLMHatasi(f"{klasor}: mesaj bulunamadı (mesajlar/ klasöründe .txt/.docx/.pdf/.eml).")
    harita = takma_adlar(mesajlar, terimler or [])
    print(f"[OK] {len(mesajlar)} mesaj · kurum bilgisi {'var' if kurum else 'yok'} · terimce {len(terimce)} terim")
    llm.onay_al(f"{len(mesajlar)} hasta mesajı gönderilecek. Mesajlar SAĞLIK VERİSİ (özel nitelikli kişisel veri) içerir; hasta adları takma adlı, "
                "telefon, e-posta, pasaport no, TCKN ve IBAN maskeli. Hastanın açık rızası ve kurum politikası sizin sorumluluğunuzdadır; "
                "mümkünse yerel model (Ollama) kullanın.", evet)
    model_yaz(mesajlar, kurum, terimce, harita)
    md = rapor_yaz(cikti, mesajlar, harita, atlanan)
    return {"mesajlar": mesajlar, "harita": harita, "kurum": kurum, "terimce": terimce, "md": md}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Hasta mesajlarını tıbbi terimleri koruyarak çevirir ve hastanın dilinde cevap taslağı hazırlar.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "ornek_gelen_kutusu",
                   help="Klasör: mesajlar/ (.txt/.docx/.pdf/.eml), kurum_bilgileri.txt, terimce.csv")
    p.add_argument("--gizle", nargs="*", default=[], metavar="AD", help="Ayrıca maskelenecek adlar (refakatçi, doktor vb.)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "hasta_mesajlari.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.cikti, a.gizle, a.evet)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    for m in sorted(s["mesajlar"], key=lambda m: SIRA[m.yanit["aciliyet"]]):
        if m.yanit["aciliyet"] == "Acil":
            print(f"[X] ACİL {m.no}: {geri_ac(m.yanit['aciliyet_gerekcesi'], s['harita'])}")
        for o, ac in m.sorunlar:
            if o == "yüksek":
                print(f"[!] {m.no}: {ac}")
    print(f"[OK] {a.cikti.resolve()} · cevap taslakları: {s['md'].resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
