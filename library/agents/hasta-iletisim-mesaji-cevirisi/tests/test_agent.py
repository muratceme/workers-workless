"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

DOSYA = KLASOR / "ornek_veri" / "ornek_gelen_kutusu"
MESAJLAR = []


def yanit(dil, dil_adi, ceviri, aciliyet, cevap, cevap_tr, terimler=(), yer=(), doktor=""):
    return {"dil": dil, "dil_adi": dil_adi, "ceviri_tr": ceviri, "terimler": [{"kaynak": k, "turkce": t} for k, t in terimler],
            "talepler": ["Bilgi talebi"], "aciliyet": aciliyet, "aciliyet_gerekcesi": "Gerekçe.", "cevap": cevap, "cevap_tr": cevap_tr,
            "koordinatore_sorular": list(yer), "doktora_iletilecek": doktor}


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    kullanici = kullanici.split("</mesaj>")[0]        # terimce tüm mesajlarda aynı; yalnız mesaj gövdesine bak
    if "knee replacement" in kullanici:      # doğru çeviri, yer tutuculu cevap
        return yanit("en", "İngilizce",
                     "64 yaşındayım; MR'ım ciddi osteoartrit gösteriyor, total diz protezi önerildi. Tip 2 diyabetim var, günde iki kez metformin "
                     "500 mg ve günde bir kez varfarin 5 mg kullanıyorum.", "Öncelikli",
                     "Dear [HASTA-1], thank you for your message. Airport transfer is free of charge. Your length of stay will be [KALIŞ SÜRESİ] "
                     "after our doctor's evaluation.",
                     "Sayın [HASTA-1], mesajınız için teşekkürler. Havalimanı transferi ücretsizdir. Kalış süreniz doktor değerlendirmesinden sonra "
                     "[KALIŞ SÜRESİ] olacaktır.", [("Warfarin", "varfarin")], ["[KALIŞ SÜRESİ]: doktor değerlendirmesi sonrası"],
                     "Varfarin kullanımı: ameliyat öncesi plan.")
    if "Zahnimplantate" in kullanici:        # doz eksik çeviri, uydurma fiyat
        return yanit("de", "Almanca", "Alt çeneye iki implant istiyorum. Osteoporozum var ve haftada bir alendronat kullanıyorum.", "Rutin",
                     "Sehr geehrte Frau [HASTA-2], die Behandlung kostet etwa 1500 Euro.", "Sayın [HASTA-2], tedavi yaklaşık 1500 Euro tutar.")
    if "hip surgery" in kullanici:           # kırmızı bayrak, aciliyet yanlış
        return yanit("en", "İngilizce", "28.09.2026 tarihinde ameliyat oldum. Yaram kızarık ve şiş, sarı akıntı var, ateşim 38.9 °C. Yürürken nefes "
                     "darlığı hissediyorum.", "Öncelikli", "Dear [HASTA-3], our doctor will contact you.", "Sayın [HASTA-3], doktorumuz sizinle "
                     "iletişime geçecek.")
    return yanit("en", "İngilizce", "Kardiyolog muayenesi istiyorum. Günde bir kez bisoprolol 5 mg kullanıyorum. 20 Ekim için randevu "
                 "alabilir miyim?", "Rutin", "Здравствуйте, [HASTA-4]! Запись на [TARİH] возможна.", "Merhaba [HASTA-4]! [TARİH] için kayıt mümkün.",
                 yer=["[TARİH]: uygun randevu"])


class GirdiTesti(unittest.TestCase):
    def test_okuma_ve_yardimcilar(self):
        ms, kurum, terimce, atlanan = agent.dosyalari_oku(DOSYA)
        self.assertEqual([(m.no, m.hasta[0], m.tarih) for m in ms][:2], [("M-001", "John Sample", "06.10.2026 09:12"), ("M-002", "Anna Beispiel", "06.10.2026 14:40")])
        self.assertIn("transfer", kurum)
        self.assertEqual(len(terimce), 14)
        self.assertEqual(atlanan, [])
        self.assertEqual(agent.olculer("38,9 °C, 5 мг, 500mg, 10 ml, 12 kg"), [("38.9", "°C"), ("5", "mg"), ("500", "mg"), ("10", "ml"), ("12", "kg")])
        self.assertEqual(agent.rakamlar("[TARİH] 38,9 ve 7/24"), {"38.9", "7", "24"})

    def test_maskeleme(self):
        ms, *_ = agent.dosyalari_oku(DOSYA)
        h = agent.takma_adlar(ms, [])
        m4 = agent.maskele(ms[3].metin, h)
        self.assertIn("[PASAPORT]", m4)
        self.assertNotIn("1234567", m4)
        self.assertTrue(m4.rstrip().endswith("[HASTA-4]"))
        self.assertIn("[TELEFON]", agent.maskele(ms[0].metin, h))
        self.assertEqual(agent.geri_ac("Sayın [HASTA-4]", h), "Sayın Ivan Primer")


class AgentTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "h.xlsx"
        MESAJLAR.clear()
        with mock.patch.object(agent.llm, "json_iste", sahte_json_iste):
            cls.s = agent.calistir(DOSYA, cls.cikti, evet=True)
        cls.m = {m.no: m for m in cls.s["mesajlar"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def sorun(self, no, onem="yüksek"):
        return [a for o, a in self.m[no].sorunlar if o == onem]

    def test_gonderilen(self):
        self.assertEqual(len(MESAJLAR), 4)
        for gizli in ("John", "Sample", "Иван", "7700", "751234567", "75 1234567"):
            self.assertFalse(any(gizli in x for x in MESAJLAR), gizli)

    def test_dogru_mesaj(self):
        self.assertEqual(self.sorun("M-001"), [])
        self.assertEqual(self.m["M-001"].yanit["yer_tutucular"], ["[KALIŞ SÜRESİ]"])

    def test_doz_ve_uydurma_sayi(self):
        s = self.sorun("M-002")
        self.assertTrue(any("70 mg" in a for a in s))
        self.assertTrue(any("1500" in a and "olmayan sayı" in a for a in s))
        self.assertTrue(any("Terimce: 'Zahnimplantat'" in a for a in self.sorun("M-002", "orta")))

    def test_kirmizi_bayrak(self):
        s = self.sorun("M-003")
        self.assertTrue(any("kırmızı bayrak" in a and "ates" in a for a in s))

    def test_alfabe_ve_terimce(self):
        orta = self.sorun("M-004", "orta")
        self.assertTrue(any("Kiril" in a for a in orta))
        self.assertFalse(any("Terimce" in a for a in orta))
        self.assertTrue(any("Terimce: 'Zahnimplantat'" in a for a in self.sorun("M-002", "orta")))

    def test_ciktilar(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Mesajlar", "Terimler", "Kontroller"])
        ms = wb["Mesajlar"]
        self.assertEqual([ms.cell(r, 1).value for r in range(2, 6)], ["M-001", "M-003", "M-002", "M-004"])
        self.assertIn("Dear John Sample", ms["I2"].value)
        md = (self.s["md"] / "M-004.md").read_text(encoding="utf-8")
        self.assertIn("Здравствуйте, Ivan Primer!", md)
        self.assertIn("Taslaktır", md)


class CliTesti(unittest.TestCase):
    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(agent.main(["--girdi", str(Path(tmp) / "yok")]), 1)
            self.assertEqual(agent.main(["--girdi", tmp]), 1)
            with mock.patch.object(agent.llm, "json_iste", sahte_json_iste):
                self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "x.xlsx")]), 0)


if __name__ == "__main__":
    unittest.main()
