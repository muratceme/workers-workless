"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import json
import os
import re
import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
CEVAP = json.loads((Path(__file__).parent / "sahte_cevaplar.json").read_text(encoding="utf-8"))
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    """Satış Müdürü'nde ilk turda 'en az 8 yıl' deneyime '10+ yıl' düzeyini verir."""
    MESAJLAR.append(kullanici)
    ad = re.search(r'<is_tanimi pozisyon="([^"]+)"', kullanici)[1]
    return copy.deepcopy(CEVAP[f"{ad}_hatali" if ad == "Satış Müdürü" and "<duzeltme>" not in kullanici else ad])


def hazir():
    return (agent.plan_oku(ORNEK / "degerleme_plani.csv"), agent.kademe_oku(ORNEK / "kademeler.csv"),
            {t["pozisyon"]: t for t in agent.tanimlari_oku(ORNEK / "is_tanimlari.txt")})


class KodTesti(unittest.TestCase):
    def test_okuma(self):
        plan, kad, t = hazir()
        self.assertEqual(list(plan), ["Eğitim", "Deneyim (yıl)", "Yönetim (bağlı çalışan)", "Bütçe sorumluluğu (TL)", "Karar yetkisi", "İletişim"])
        self.assertEqual([f["sayisal"] for f in plan.values()], [False, True, True, True, False, False])
        self.assertEqual(plan["Bütçe sorumluluğu (TL)"]["duzeyler"]["10 milyon TL'ye kadar"]["ust"], D(10000000))
        self.assertEqual(kad[-1], ("K5", D(361), None))
        self.assertEqual((t["Muhasebe Uzmanı"]["departman"], t["Muhasebe Uzmanı"]["mevcut"]), ("Finans", "K3"))
        self.assertEqual(agent.kanit_sayilari("Yıllık 12 milyon TL ve 1,5 milyar; 500 bin TL; 4 bölge"), [D(12000000), D(1500000000), D(500000), D(4)])
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "t.txt"
            y.write_text("Pozisyon: A\nmetin\nPozisyon: A\nmetin\n", encoding="utf-8")
            with self.assertRaises(llm.LLMHatasi):
                agent.tanimlari_oku(y)

    def test_denetim(self):
        plan, _, t = hazir()
        d = lambda a, c: agent.denetle(t[a], plan, c)
        self.assertEqual([d(a, CEVAP[a]) for a in t], [[], [], []])
        self.assertEqual(d("Satış Müdürü", CEVAP["Satış Müdürü_hatali"]),
                         ["Satış Müdürü: Deneyim (yıl): kanıttaki değer 8 '10+ yıl' aralığında değil (10 ve üstü); '5–9 yıl' olmalı"])
        c = copy.deepcopy(CEVAP["Muhasebe Uzmanı"])
        c["faktorler"][0].update(duzey="Lisans üstü")
        c["faktorler"][2].update(duzey="1–5 kişi", kanit="uzman yardımcısına iş dağılımında rehberlik eder")     # rehberlik, sayı yok
        c["faktorler"][4].update(kanit="Kayıt kararlarını kendisi verir")                                       # iş tanımında yok
        c["faktorler"][5].update(duzey="Dış müzakere", kanit="SMMM belgesi tercih sebebi")
        c["faktorler"].pop(3)
        c["faktorler"].append(copy.deepcopy(c["faktorler"][0]) | {"faktor": "Fiziksel efor"})
        s = d("Muhasebe Uzmanı", c)
        for p in ("değerlendirilmeyen faktör: Bütçe", "planda olmayan ya da tekrar eden faktör: Fiziksel efor", "'Lisans üstü' planda yok",
                  "kanıtta sayı yok", "Karar yetkisi: kanıt iş tanımında aynen geçmiyor", "tercih edilen bir niteliğe"):
            self.assertTrue(any(p in x for x in s), p)
        c = copy.deepcopy(CEVAP["Satış Müdürü"])
        c["faktorler"][3].update(duzey="10 milyon TL'ye kadar")
        self.assertIn("'10 milyon TL üstü' olmalı", d("Satış Müdürü", c)[0])

    def test_puanlama(self):
        plan, kad, t = hazir()
        p = {a: agent.puanla(t[a], plan, kad, CEVAP[a]) for a in t}
        self.assertEqual([(x["toplam"], x["kademe"], x["durum"]) for x in p.values()],
                         [(D(160), "K2", "Uyumlu"), (D(180), "K2", "İnceleme"), (D(420), "K5", "Uyumlu")])
        c = copy.deepcopy(CEVAP["Ekip Lideri"])
        c["faktorler"][0]["duzey"] = ""
        x = agent.puanla(t["Ekip Lideri"], plan, kad, c)
        self.assertEqual((x["kademe"], x["durum"], x["eksik"]), ("", "Eksik faktör", ["Eğitim"]))
        c["faktorler"][0]["duzey"] = "Yüksek lisans"                   # 160 - 20 + 80 = 220 → K3, alt sınır 201'e 19 puan
        x = agent.puanla(t["Ekip Lideri"] | {"mevcut": ""}, plan, kad, c, sinir=20)
        self.assertEqual((x["toplam"], x["kademe"], x["durum"]), (D(220), "K3", "Yeni"))
        self.assertIn("alt sınırına 19 puan", x["notlar"][0])


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ornek_ve_duzeltme(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "d.xlsx"
            self.assertEqual(agent.main(["--evet", "--cikti", str(cikti)]), 0)
            self.assertEqual(len(MESAJLAR), 4)                          # 3 pozisyon + Satış Müdürü düzeltmesi
            self.assertIn("10 milyon TL üstü [10.000.001 ve üstü]", MESAJLAR[0])
            wb = load_workbook(cikti)
            r = [[c.value for c in x] for x in wb["Pozisyonlar"].iter_rows(min_row=2)]
            self.assertEqual([x[9:12] for x in r], [[160, "K2", "Uyumlu"], [180, "K2", "İnceleme"], [420, "K5", "Uyumlu"]])
            self.assertEqual(r[2][4], 60)                               # deneyim düzeltildi
            f = [[c.value for c in x] for x in wb["Faktörler"].iter_rows(min_row=2)]
            self.assertEqual(len(f), 18)
            self.assertEqual(wb["Kontrol"]["A2"].value, "Son kontrol: Kod denetiminden sorunsuz geçti")

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(agent.main(["--plan", str(Path(tmp) / "yok.csv"), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
