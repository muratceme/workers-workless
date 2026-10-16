"""
İş Tanımından İş Değerleme Puanlaması — Workers / Workless AI Agent
İnsan Kaynakları › Ücretlendirme ve Yan Haklar Uzmanı

Serbest metin iş tanımlarını şirketin puan faktör planına göre değerlendirir; puanı ve kademeyi kod hesaplar.
1. Kod iş tanımlarını ("Pozisyon: Ad | Departman: … | Mevcut Kademe: …" satırıyla başlayan bloklar), değerleme planını
   (Faktör; Düzey; Tanım; Alt; Üst; Puan) ve kademe tablosunu (Kademe; Alt Puan; Üst Puan) okur.
2. Model her faktör için plandaki düzeylerden birini seçer ve iş tanımından kanıt cümlesini aynen yazar.
3. Kod denetler: her faktör bir kez mi; düzey planda var mı; kanıt iş tanımında aynen geçiyor mu; sayısal faktörlerde
   (deneyim yılı, bağlı çalışan, bütçe) kanıttaki sayı seçilen düzeyin aralığında mı; kanıt yalnız "tercihen" yazan bir
   niteliğe mi dayanıyor. Sorun varsa yalnız o pozisyon modele geri gönderilir.
4. Kod puanı toplar, kademeyi bulur, mevcut kademeyle karşılaştırır ve sınıra yakın puanları işaretler.

Kullanım:
    python agent.py                                                          # örnek 3 pozisyon
    python agent.py --tanimlar is_tanimlari.docx --plan plan.xlsx --kademeler kademeler.xlsx
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}
CARPAN = {"milyar": Decimal(10) ** 9, "milyon": Decimal(10) ** 6, "bin": Decimal(1000)}
TERCIH = re.compile(r"\b(tercih|avantaj|tercihen)\w*")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def sayi(x) -> Decimal | None:
    s = re.sub(r"[^\d,.\-]", "", str(x if x is not None else ""))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1 or re.fullmatch(r"\d{1,3}\.\d{3}", s):
        s = s.replace(".", "")
    try:
        return Decimal(s) if s not in ("", "-", ".") else None
    except InvalidOperation:
        return None


def kanit_sayilari(metin: str) -> list[Decimal]:
    """Kanıttaki sayılar; '12 milyon', '1,5 milyar', '500 bin' çarpanlarıyla."""
    out = []
    for m in re.finditer(r"(\d+(?:[.,]\d+)*)\s*(milyar|milyon|bin)?\b", str(metin).replace("İ", "i").lower().translate(_TR)):
        v = sayi(m[1])
        if v is not None:
            out.append(v * CARPAN.get(m[2] or "", 1))
    return out


# ----------------------------------------------------------------------------
# Girdi
# ----------------------------------------------------------------------------

def tablo_oku(yol: Path) -> list[dict]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, read_only=True, data_only=True)
        s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=metin.split("\n")[0].count)))
    s = [r for r in s if any(c not in (None, "") for c in r)]
    bas = [katla(c) for c in s[0]] if s else []
    return [dict(zip(bas, list(r) + [None] * (len(bas) - len(r)))) for r in s[1:]]


def plan_oku(yol: Path) -> dict[str, dict]:
    plan = {}
    for r in tablo_oku(yol):
        f, d, p = str(r.get("faktor") or "").strip(), str(r.get("duzey") or "").strip(), sayi(r.get("puan"))
        if not f or not d:
            continue
        if p is None:
            raise llm.LLMHatasi(f"{yol.name}: '{f} / {d}' için puan okunamadı.")
        alt, ust = sayi(r.get("alt")), sayi(r.get("ust"))
        x = plan.setdefault(f, {"ad": f, "duzeyler": {}, "sayisal": False})
        x["duzeyler"][d] = {"ad": d, "tanim": str(r.get("tanim") or "").strip(), "alt": alt, "ust": ust, "puan": p}
        x["sayisal"] |= alt is not None or ust is not None
    if not plan:
        raise llm.LLMHatasi(f"{yol.name}: 'Faktör; Düzey; Tanım; Alt; Üst; Puan' sütunlarıyla en az bir satır gerekli.")
    return plan


def kademe_oku(yol: Path) -> list[tuple[str, Decimal, Decimal | None]]:
    k = [(str(r.get("kademe") or "").strip(), sayi(r.get("alt puan")), sayi(r.get("ust puan"))) for r in tablo_oku(yol)]
    k = [x for x in k if x[0] and x[1] is not None]
    if not k:
        raise llm.LLMHatasi(f"{yol.name}: 'Kademe; Alt Puan; Üst Puan' sütunları gerekli.")
    return sorted(k, key=lambda x: x[1])


def tanimlari_oku(yol: Path) -> list[dict]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: iş tanımı dosyası bulunamadı.")
    out, son = [], None
    for s in belge.metin_oku(yol).splitlines():
        m = re.match(r"^\s*Pozisyon\s*:\s*(.+)$", s, re.I)
        if m:
            alan = [x.strip() for x in m[1].split("|")]
            ek = {katla(a.split(":", 1)[0]): a.split(":", 1)[1].strip() for a in alan[1:] if ":" in a}
            son = {"pozisyon": alan[0], "departman": ek.get("departman", ""), "mevcut": ek.get("mevcut kademe", ek.get("kademe", "")), "metin": ""}
            out.append(son)
        elif son and s.strip():
            son["metin"] = (son["metin"] + "\n" + llm.maskele(s.strip())).strip()
    out = [x for x in out if x["metin"]]
    if not out:
        raise llm.LLMHatasi(f"{yol.name}: 'Pozisyon: Ad | Departman: … | Mevcut Kademe: …' satırıyla başlayan iş tanımı yok.")
    tekrar = sorted({x["pozisyon"] for x in out if [y["pozisyon"] for y in out].count(x["pozisyon"]) > 1})
    if tekrar:
        raise llm.LLMHatasi(f"Aynı pozisyon birden çok kez geçiyor: {', '.join(tekrar)}")
    return out


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

SEMA = {"type": "object", "properties": {
    "faktorler": {"type": "array", "items": {"type": "object", "properties": {"faktor": S, "duzey": S, "kanit": S, "gerekce": S},
                                             "required": ["faktor", "duzey", "kanit", "gerekce"], "additionalProperties": False}},
    "belirsizlikler": {"type": "array", "items": S}},
    "required": ["faktorler", "belirsizlikler"], "additionalProperties": False}


def aralik(d: dict) -> str:
    a, u = d["alt"], d["ust"]
    return f"{a:,.0f}–{u:,.0f}".replace(",", ".") if u is not None else f"{a:,.0f} ve üstü".replace(",", ".")


def plan_metni(plan: dict) -> str:
    l = []
    for f in plan.values():
        l.append(f"Faktör: {f['ad']}" + (" (sayısal)" if f["sayisal"] else ""))
        for d in f["duzeyler"].values():
            l.append(f"  - {d['ad']}" + (f" [{aralik(d)}]" if f["sayisal"] else "") + (f": {d['tanim']}" if d["tanim"] else ""))
    return "\n".join(l)


def uygun_duzey(f: dict, v: Decimal) -> str | None:
    return next((d["ad"] for d in f["duzeyler"].values() if d["alt"] is not None and v >= d["alt"] and (d["ust"] is None or v <= d["ust"])), None)


def denetle(t: dict, plan: dict, c: dict) -> list[str]:
    ad, km, sorun = t["pozisyon"], katla(t["metin"]), []
    adlar = [x["faktor"] for x in c["faktorler"]]
    eksik = [f for f in plan if f not in adlar]
    if eksik:
        sorun.append(f"{ad}: değerlendirilmeyen faktör: {', '.join(eksik)} (dayanak yoksa düzeyi boş bırakıp yine de yazın)")
    fazla = sorted({f for f in adlar if f not in plan} | {f for f in adlar if adlar.count(f) > 1})
    if fazla:
        sorun.append(f"{ad}: planda olmayan ya da tekrar eden faktör: {', '.join(fazla)}")
    for x in c["faktorler"]:
        f = plan.get(x["faktor"])
        if not f or not x["duzey"]:
            continue
        on = f"{ad}: {x['faktor']}"
        d = f["duzeyler"].get(x["duzey"])
        if not d:
            sorun.append(f"{on}: '{x['duzey']}' planda yok; şunlardan biri olmalı: {', '.join(f['duzeyler'])}")
            continue
        kk = katla(x["kanit"])
        if len(kk.split()) < 2 or kk not in km:
            sorun.append(f"{on}: kanıt iş tanımında aynen geçmiyor; ilgili cümleyi kelimesi kelimesine yazın")
            continue
        parca = [p for p in re.split(r"[,;.]", x["kanit"]) if katla(p)]
        if parca and all(TERCIH.search(katla(p)) for p in parca):
            sorun.append(f"{on}: kanıt yalnız tercih edilen bir niteliğe dayanıyor; aranan asgari şartı esas alın")
        if f["sayisal"]:
            say = kanit_sayilari(x["kanit"])
            if not say:
                if d["alt"] != 0:
                    sorun.append(f"{on}: kanıtta sayı yok ama '{x['duzey']}' düzeyi sayı gerektiriyor ({aralik(d)})")
            else:
                v = max(say)
                if not (v >= d["alt"] and (d["ust"] is None or v <= d["ust"])):
                    dogru = uygun_duzey(f, v)
                    sorun.append(f"{on}: kanıttaki değer {v:,.0f}".replace(",", ".") + f" '{x['duzey']}' aralığında değil ({aralik(d)})"
                                 + (f"; '{dogru}' olmalı" if dogru else ""))
    return sorun


def uret(tanimlar: list[dict], plan: dict, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    pm = plan_metni(plan)
    sonuc = dict(onceki or {})
    hedef = tanimlar if duzeltme is None else [t for t in tanimlar if any(s.startswith(t["pozisyon"] + ":") for s in duzeltme)]
    for t in hedef:
        mesaj = f"<plan>\n{pm}\n</plan>\n<is_tanimi pozisyon=\"{t['pozisyon']}\" departman=\"{t['departman']}\">\n{t['metin']}\n</is_tanimi>"
        d = [s.split(": ", 1)[1] for s in (duzeltme or []) if s.startswith(t["pozisyon"] + ":")]
        if d:
            mesaj += ("\n<duzeltme>\nÖnceki cevap: " + json.dumps(sonuc.get(t["pozisyon"]), ensure_ascii=False) + "\nKodun bulduğu sorunlar:\n"
                      + "\n".join(f"- {s}" for s in d) + "\n</duzeltme>")
        sonuc[t["pozisyon"]] = llm.json_iste(sistem, mesaj, SEMA)
    return sonuc


def duzeltme_dongusu(uret_f, denetle_f, tur: int) -> tuple[dict, list[str], int]:
    taslak = uret_f(None, None)
    kalan = denetle_f(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret_f(kalan, taslak)
        kalan = denetle_f(taslak)
    return taslak, kalan, yapilan


# ----------------------------------------------------------------------------
# Hesap ve çıktı
# ----------------------------------------------------------------------------

def puanla(t: dict, plan: dict, kademeler: list, c: dict, sinir: int = 10) -> dict:
    sec = {x["faktor"]: x for x in c["faktorler"] if x["faktor"] in plan}
    puanlar = {f: plan[f]["duzeyler"][sec[f]["duzey"]]["puan"] for f in plan if f in sec and sec[f]["duzey"] in plan[f]["duzeyler"]}
    eksik = [f for f in plan if f not in puanlar]
    toplam = sum(puanlar.values(), Decimal(0))
    kademe = next((k for k, a, u in kademeler if toplam >= a and (u is None or toplam <= u)), "") if not eksik else ""
    notlar = []
    if eksik:
        durum = "Eksik faktör"
        notlar.append(f"Düzeyi belirlenemeyen faktör: {', '.join(eksik)}; kademe hesaplanmadı.")
    elif not kademe:
        durum = "Kademe bulunamadı"
        notlar.append(f"{toplam:g} puan kademe tablosunun hiçbir aralığına düşmüyor.")
    elif t["mevcut"] and katla(t["mevcut"]) != katla(kademe):
        durum = "İnceleme"
        notlar.append(f"Puan {toplam:g} → {kademe}; mevcut kademe {t['mevcut']}.")
    else:
        durum = "Uyumlu" if t["mevcut"] else "Yeni"
    if kademe:
        for k, a, u in kademeler:
            if k == kademe:
                yakin = [(b, ad) for b, ad in ((a, "alt"), (u, "üst")) if b is not None and abs(toplam - b) <= sinir and not (ad == "alt" and a == 0)]
                if yakin:
                    notlar.append(f"Kademe {yakin[0][1]} sınırına {abs(toplam - yakin[0][0]):g} puan uzaklıkta.")
    return {"puanlar": puanlar, "toplam": toplam, "kademe": kademe, "durum": durum, "notlar": notlar, "eksik": eksik}


def tablo(ws, bas: list[str], satirlar: list[list], ai: set[int] = frozenset(), genis: dict | None = None) -> None:
    ws.append(bas)
    for h in ws[1]:
        h.fill, h.font, h.alignment = BASLIK_DOLGU, BASLIK_YAZI, UST
    for r in satirlar:
        ws.append(r)
        n = ws.max_row
        for j, v in enumerate(r, 1):
            ws.cell(n, j).alignment = UST
            if j in ai and v not in (None, ""):
                ws.cell(n, j).fill = AI_DOLGU
    for j in range(1, len(bas) + 1):
        ws.column_dimensions[get_column_letter(j)].width = (genis or {}).get(j, 14)
    ws.freeze_panes = "A2"


def rapor_yaz(cikti: Path, tanimlar: list[dict], plan: dict, kademeler: list, sonuc: dict, hesap: dict, kalan: list[str], tur: int) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Pozisyonlar"
    fa = list(plan)
    rr = []
    for t in tanimlar:
        h = hesap[t["pozisyon"]]
        sor = "; ".join(x.split(": ", 1)[1] for x in kalan if x.startswith(t["pozisyon"] + ":"))
        rr.append([t["pozisyon"], t["departman"], t["mevcut"]] + [float(h["puanlar"][f]) if f in h["puanlar"] else None for f in fa]
                  + [float(h["toplam"]), h["kademe"], h["durum"], " ".join(h["notlar"]), sor or "Geçti"])
    tablo(ws, ["Pozisyon", "Departman", "Mevcut Kademe"] + fa + ["Toplam Puan", "Önerilen Kademe", "Durum", "Not", "Son Kontrol"], rr,
          genis={1: 22, len(fa) + 7: 46, len(fa) + 8: 40})
    renk = {"Uyumlu": YESIL, "Yeni": YESIL, "İnceleme": SARI}
    for n, r in enumerate(rr, 2):
        ws.cell(n, len(fa) + 6).fill = renk.get(r[len(fa) + 5], KIRMIZI)
        ws.cell(n, len(fa) + 8).fill = YESIL if r[-1] == "Geçti" else KIRMIZI
    g = wb.create_sheet("Faktörler")
    rr = []
    for t in tanimlar:
        sec = {x["faktor"]: x for x in sonuc[t["pozisyon"]]["faktorler"]}
        for f in fa:
            x = sec.get(f, {"duzey": "", "kanit": "", "gerekce": ""})
            p = hesap[t["pozisyon"]]["puanlar"].get(f)
            rr.append([t["pozisyon"], f, x["duzey"] or "Belirlenemedi", float(p) if p is not None else None, x["kanit"], x["gerekce"]])
    tablo(g, ["Pozisyon", "Faktör", "Düzey", "Puan", "İş Tanımındaki Kanıt", "Gerekçe"], rr, ai={3, 5, 6}, genis={1: 22, 2: 24, 3: 22, 5: 60, 6: 50})
    g = wb.create_sheet("Açık Noktalar")
    tablo(g, ["Pozisyon", "Açıklama"], [[t["pozisyon"], b] for t in tanimlar for b in sonuc[t["pozisyon"]]["belirsizlikler"]], ai={2}, genis={1: 22, 2: 100})
    g = wb.create_sheet("Değerleme Planı")
    tablo(g, ["Faktör", "Düzey", "Tanım", "Alt", "Üst", "Puan"],
          [[f["ad"], d["ad"], d["tanim"], float(d["alt"]) if d["alt"] is not None else None, float(d["ust"]) if d["ust"] is not None else None, float(d["puan"])]
           for f in plan.values() for d in f["duzeyler"].values()], genis={1: 24, 2: 24, 3: 60})
    tablo(g, [], [[], ["Kademe", "Alt Puan", "Üst Puan"]] + [[k, float(a), float(u) if u is not None else None] for k, a, u in kademeler])
    k = wb.create_sheet("Kontrol")
    k.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    k.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    k.cell(2, 1).fill = KIRMIZI if kalan else YESIL
    k.append(["Mor hücreler modelin seçtiği düzey ve kanıtlardır; puan, kademe ve karşılaştırma koddur. Değerleme sonucu iş değerleme "
              "komitesinin onayı olmadan ücret kararına esas alınmamalıdır. Önerilen kademe, 'Ücret Bandı Karşılaştırması' kod bloğunun "
              "çalışan dosyasındaki Kademe sütununa işlenerek kullanılabilir."])
    k.column_dimensions["A"].width = 140
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(t_yol: Path, p_yol: Path, k_yol: Path, cikti: Path, tur: int = 2, sinir: int = 10, evet: bool = False) -> dict:
    plan, kademeler, tanimlar = plan_oku(p_yol), kademe_oku(k_yol), tanimlari_oku(t_yol)
    print(f"[OK] {len(tanimlar)} iş tanımı · {len(plan)} faktör · {len(kademeler)} kademe")
    llm.onay_al(f"{len(tanimlar)} iş tanımı ve değerleme planı gönderilecek (telefon, e-posta, TCKN maskeli).", evet)
    sonuc, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(tanimlar, plan, d, o), lambda x: [s for t in tanimlar for s in denetle(t, plan, x[t["pozisyon"]])], tur)
    hesap = {t["pozisyon"]: puanla(t, plan, kademeler, sonuc[t["pozisyon"]], sinir) for t in tanimlar}
    rapor_yaz(cikti, tanimlar, plan, kademeler, sonuc, hesap, kalan, yapilan)
    return {"tanimlar": tanimlar, "sonuc": sonuc, "hesap": hesap, "kalan": kalan, "tur": yapilan}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="İş tanımlarını puan faktör planına göre değerlendirir; puanı ve kademeyi kod hesaplar.")
    p.add_argument("--tanimlar", type=Path, default=ORNEK / "is_tanimlari.txt", help="İş tanımları (.txt/.docx/.pdf); her blok 'Pozisyon: Ad | Departman: … | Mevcut Kademe: …' satırıyla başlar")
    p.add_argument("--plan", type=Path, default=ORNEK / "degerleme_plani.csv", help="Faktör; Düzey; Tanım; Alt; Üst; Puan (.xlsx/.csv)")
    p.add_argument("--kademeler", type=Path, default=ORNEK / "kademeler.csv", help="Kademe; Alt Puan; Üst Puan (.xlsx/.csv)")
    p.add_argument("--sinir", type=int, default=10, help="Kademe sınırına bu kadar puan yakınsa işaretle (varsayılan 10)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "is_degerleme.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.tanimlar, a.plan, a.kademeler, a.cikti, a.tur, a.sinir, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    for t in s["tanimlar"]:
        h = s["hesap"][t["pozisyon"]]
        print(f"[OK] {t['pozisyon']} · {h['toplam']:g} puan · {h['kademe'] or '-'} · {h['durum']}" + (f" ({' '.join(h['notlar'])})" if h["notlar"] else ""))
    for x in s["kalan"]:
        print(f"[!] Kalan: {x}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
