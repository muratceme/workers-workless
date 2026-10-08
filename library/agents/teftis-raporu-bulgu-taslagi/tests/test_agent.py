"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

DOSYA = KLASOR / "ornek_veri" / "ornek_sube"
MESAJLAR = []


def bulgu(baslik, nolar, kriterler, durum="Durum.", risk="Orta"):
    return {"baslik": baslik, "tespit_nolari": nolar, "durum": durum, "kriterler": [{"madde": m, "alinti": a} for m, a in kriterler],
            "neden": "Kök neden birimle görüşülerek belirlenmeli.", "etki": "Risk doğurmaktadır.", "risk_duzeyi": risk,
            "risk_gerekcesi": "Hata oranı %18,2.", "oneriler": ["Eksikler tamamlanmalı."]}


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return {"bulgular": [
        bulgu("Kasa sayım tutanaklarında ikinci imza eksikliği", ["T1"],
              [("K1-md5", "Sayım sonucu kasa sayım tutanağına yazılır ve tutanak sayıma katılan iki yetkili tarafından imzalanır.")],
              "22 iş gününün 4'ünde ikinci yetkili imzası olmadığı görülmüştür."),
        bulgu("Talimatlarda imza kontrolü ibaresi eksikliği", ["T2"],
              [("K1-md8", "Personel talimata imza kontrol edildi ibaresini yazar.")],
              "[PERSONEL-1] tarafından yapılan 3 işlemde toplam 190.000 TL tutarında talimatta ibare yoktur."),
        bulgu("Dört göz ilkesine aykırı işlemler", ["T4", "T5", "T9"],
              [("K1-md14", "Mesai saatleri dışında yapılması zorunlu işlemler şube müdürü veya operasyon yöneticisinin yazılı onayıyla yapılır."),
               ("K9-md3", "Personel kendi hesabında işlem yapamaz.")], risk="Yüksek"),
        bulgu("Kart teslim formlarında imza eksikliği", ["T3"], []),
    ], "genel_not": ""}


class GirdiTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.t, cls.m, cls.k, cls.a = agent.dosyalari_oku(DOSYA)

    def test_okuma(self):
        self.assertEqual([t.no for t in self.t], ["T1", "T2", "T3", "T4", "T5", "T6"])
        t2 = self.t[1]
        self.assertEqual((t2.orneklem, t2.hatali, t2.tutar, t2.musteri), (D(40), D(3), D("186400.00"), ["Kurgu Tekstil Ltd. Şti.", "Ali Sınamalı"]))
        self.assertAlmostEqual(float(self.t[0].hata_orani), 4 / 22)
        self.assertEqual([m.id for m in self.m], ["K1-md1", "K1-md5", "K1-md8", "K1-md11", "K1-md14", "K1-md15", "K1-md19", "K1-md21"])
        self.assertIn("şube müdürü yazılı olarak belirler", self.m[1].metin)

    def test_madde_secimi(self):
        with mock.patch.object(agent, "TUM_MADDE_SINIRI", 0):
            secilen = [m.id for m in agent.madde_sec(self.t, self.m)]
        self.assertNotIn("K1-md1", secilen)
        for gerekli in ("K1-md5", "K1-md8", "K1-md11", "K1-md14", "K1-md15", "K1-md19"):
            self.assertIn(gerekli, secilen)

    def test_kontroller(self):
        self.assertEqual(agent.kontrol_et(self.t, self.m), [])
        t = [agent.Tespit("A", "Şube", "Kasa", "x", D(5), D(7), None, "", [], [], "", ""),
             agent.Tespit("A", "Şube", "Kasa", "y", None, D(1), None, "K", [], [], "", "")]
        metin = " ".join(a for _, a, _ in agent.kontrol_et(t, []))
        for parca in ("Mevzuat / iç düzenleme metni verilmedi", "Tespit numarası tekrar ediyor: A", "A: kanıt", "hatalı adet (7) örneklemden (5)",
                      "örneklem büyüklüğü yazılmamış", "tek bulguda birleştirilebilir"):
            self.assertIn(parca, metin)

    def test_maddesiz_metin(self):
        baslik, m = agent.maddelere_bol("Kart Prosedürü\nKartlar imza karşılığı teslim edilir.", "K2")
        self.assertEqual((baslik, [x.id for x in m]), ("Kart Prosedürü", ["K2-tum"]))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_uctan_uca(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "b.xlsx"
            s = agent.calistir(DOSYA, cikti, evet=True)
            m = MESAJLAR[0]
            for gizli in ("Mehmet Deneme", "Zeynep Örnek", "Kurgu Tekstil", "Ali Sınamalı", "Ayşe Kurgu"):
                self.assertNotIn(gizli, m)
            for var in ("[PERSONEL-1]", "[MÜŞTERİ-1]", '<madde id="K1-md5"', "hata oranı %18,2"):
                self.assertIn(var, m)
            y = s["yanit"]
            kr = [k["durum"] for b in y["bulgular"] for k in b["kriterler"]]
            self.assertEqual(kr, ["Doğrulandı", "Alıntı doğrulanamadı", "Doğrulandı", "Madde bulunamadı"])
            self.assertEqual([b["dayanak_var"] for b in y["bulgular"]], [True, False, True, False])
            self.assertEqual(y["bulgular"][2]["tespit_nolari"], ["T4", "T5"])
            sorun = " ".join(a for _, a, _ in y["sorunlar"])
            for parca in ("tespitlerde olmayan numara yok sayıldı: T9", "'K9-md3' verilen metinlerde yok", "K1-md8 alıntısı maddede birebir geçmiyor",
                          "Bulgu 4 (Kart teslim formlarında imza eksikliği): doğrulanmış dayanak yok", "T6 hiçbir bulguda kullanılmamış"):
                self.assertIn(parca, sorun)
            self.assertEqual(y["dogrulanamayan_sayilar"], ["190000"])
            md = s["md"].read_text(encoding="utf-8")
            self.assertIn("Mehmet Deneme tarafından yapılan 3 işlemde", md)
            self.assertIn("ÖRNEK BANKA A.Ş. ŞUBE OPERASYON YÖNERGESİ (KURGUSAL ÖRNEK) md. 5", md)
            self.assertIn("**Dayanak müfettişçe eklenmeli.**", md)
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Özet", "Bulgular", "Tespitler", "Dayanak Kontrolü", "Kontroller"])
            self.assertEqual((wb["Bulgular"]["K1"].value, wb["Bulgular"]["L1"].value), ("Müfettiş Onayı", "Birim Cevabı"))
            self.assertEqual(wb["Tespitler"]["J5"].value, "3")                        # T4 → bulgu 3
            self.assertEqual(wb["Tespitler"]["J7"].value, "—")                        # T6

    def test_klasor_yok(self):
        with self.assertRaises(llm.LLMHatasi):
            agent.calistir(Path("olmayan_klasor"), Path("x.xlsx"), evet=True)


if __name__ == "__main__":
    unittest.main()
