"""
Teftiş Raporu Bulgu Taslağı — Workers / Workless AI Agent
Bankacılık › Teftiş Kurulu › Müfettiş

1. Klasördeki çalışma kâğıdı okunur (Tespit No, Birim, Konu, Tespit, Örneklem, Hatalı, Tutar, Kanıt, Personel,
   Müşteri, Kök Neden, Birim Açıklaması) ve kullanıcının verdiği mevzuat / iç düzenleme metinleri (mevzuat/
   klasörü veya adında yönetmelik, yönerge, prosedür, genelge geçen .txt/.pdf/.docx) maddelere bölünür.
2. Kod hesaplar ve kontrol eder: hata oranı, kanıtı olmayan tespit, hatalı > örneklem, aynı birim ve konuda birden
   çok tespit (birleştirme adayı). Madde sayısı çoksa her tespitle ortak kelimesi en fazla olan maddeler seçilir.
3. Personel ve müşteri adları [PERSONEL-n] / [MÜŞTERİ-n] takma adlarıyla, telefon, e-posta, TCKN ve IBAN maskelenir.
   Model her bulgu için durum (tespit), kriter (yalnız verilen metinlerden, birebir alıntıyla), neden, etki, risk
   düzeyi ve öneri yazar.
4. Kod modelin çıktısını denetler: alıntı gösterilen maddede birebir geçiyor mu, madde verilen metinlerde var mı,
   her tespit bir bulguda kullanılmış mı, uydurma tespit numarası ve girdilerde olmayan sayı var mı. Dayanağı
   doğrulanamayan bulgular "dayanak müfettişçe eklenmeli" olarak işaretlenir.
5. Çıktı: Markdown bulgu taslakları (takma adlar geri açılmış) + Excel (Özet, Bulgular + 'Müfettiş Onayı' ve
   'Birim Cevabı', Tespitler, Dayanak Kontrolü, Kontroller).

Kullanım:
    python agent.py                                          # örnek: kurgusal şube teftişi
    python agent.py --girdi ./teftis_klasoru --gizle "Ek Kişi Adı"
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
MEVZUAT_ADLARI = ("mevzuat", "yonetmelik", "yonerge", "prosedur", "genelge", "teblig", "kanun", "politika", "talimat", "duzenleme")
TUM_MADDE_SINIRI = 60   # bundan az maddede tüm maddeler gönderilir; fazlasında her tespit için en ilgili maddeler seçilir
TESPIT_BASINA_MADDE = 4
SUTUNLAR = {
    "no": ("tespit no", "no", "sira no", "bulgu no"),
    "birim": ("birim", "sube", "denetlenen birim"),
    "konu": ("konu", "alan", "surec"),
    "tespit": ("tespit", "tespit aciklamasi", "gozlem", "aciklama"),
    "orneklem": ("orneklem", "orneklem adet", "incelenen", "incelenen adet"),
    "hatali": ("hatali", "hatali adet", "istisna", "uygunsuz"),
    "tutar": ("tutar", "risk tutari", "etkilenen tutar"),
    "kanit": ("kanit", "kanit referansi", "calisma kagidi ref", "referans"),
    "personel": ("personel", "ilgili personel"),
    "musteri": ("musteri", "ilgili musteri"),
    "neden": ("kok neden", "neden"),
    "birim_aciklamasi": ("birim aciklamasi", "yonetim aciklamasi", "birim gorusu"),
}


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def katla(s) -> str:
    s = kucuk(s)
    for a, b in zip("ıöüşğçâîû", "iousgcaiu"):
        s = s.replace(a, b)
    return re.sub(r"[^a-z0-9%]+", " ", s).strip()


def para(x) -> Decimal | None:
    if x in (None, "", "—", "-"):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("₺", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tl(x) -> str:
    return "—" if x is None else f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = belge.txt_oku(yol)
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def coklu(x) -> list[str]:
    return [p.strip() for p in re.split(r"[|;\n]", str(x or "")) if p.strip()]


# ----------------------------------------------------------------------------
# Girdiler
# ----------------------------------------------------------------------------

@dataclass
class Tespit:
    no: str
    birim: str
    konu: str
    metin: str
    orneklem: Decimal | None
    hatali: Decimal | None
    tutar: Decimal | None
    kanit: str
    personel: list[str]
    musteri: list[str]
    neden: str
    birim_aciklamasi: str
    bulgular: list = field(default_factory=list)

    @property
    def hata_orani(self) -> Decimal | None:
        return self.hatali / self.orneklem if self.orneklem and self.hatali is not None else None


@dataclass
class Madde:
    id: str
    kaynak: str
    no: str
    baslik: str
    metin: str


def kagit_oku(yol: Path) -> list[Tespit]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        if "tespit" in b or "tespit aciklamasi" in b or "gozlem" in b:
            break
    else:
        raise ValueError(f"{yol.name}: 'Tespit' sütunu bulunamadı")
    es = {}
    for alan, adlar in SUTUNLAR.items():
        j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
        if j is not None:
            es[alan] = j
    al = lambda r, a: r[es[a]] if a in es and es[a] < len(r) else None  # noqa: E731
    sonuc = []
    for n, r in enumerate(s[bi + 1:], 1):
        if not al(r, "tespit"):
            continue
        sonuc.append(Tespit(str(al(r, "no") or f"T{n}").strip(), str(al(r, "birim") or "").strip(), str(al(r, "konu") or "").strip(),
                            str(al(r, "tespit")).strip(), para(al(r, "orneklem")), para(al(r, "hatali")), para(al(r, "tutar")),
                            str(al(r, "kanit") or "").strip(), coklu(al(r, "personel")), coklu(al(r, "musteri")),
                            str(al(r, "neden") or "").strip(), str(al(r, "birim_aciklamasi") or "").strip()))
    return sonuc


MADDE_RE = re.compile(r"^\s*(?:MADDE|Madde)\s+(\d+[A-Za-z/]*)\s*[-–—:.]?\s*(.*)$", re.M)


def maddelere_bol(metin: str, kaynak_id: str) -> tuple[str, list[Madde]]:
    satirlar = [x for x in metin.splitlines() if x.strip()]
    baslik = (satirlar[0].strip() if satirlar else kaynak_id)[:90]
    eslesmeler = list(MADDE_RE.finditer(metin))
    if not eslesmeler:
        return baslik, [Madde(f"{kaynak_id}-tum", baslik, "", baslik, metin.strip())]
    maddeler = []
    for i, m in enumerate(eslesmeler):
        son = eslesmeler[i + 1].start() if i + 1 < len(eslesmeler) else len(metin)
        govde = metin[m.end():son].strip()
        maddeler.append(Madde(f"{kaynak_id}-md{m.group(1)}", baslik, m.group(1), m.group(2).strip(), govde))
    return baslik, maddeler


def dosyalari_oku(klasor: Path) -> tuple[list[Tespit], list[Madde], list[str], list[str]]:
    tespitler, maddeler, kaynaklar, atlanan = [], [], [], []
    yollar = sorted(p for p in klasor.rglob("*") if p.is_file())
    for yol in yollar:
        ad = katla(yol.stem) + " " + katla(yol.parent.name if yol.parent != klasor else "")
        if yol.suffix.lower() in {".csv", ".xlsx", ".xlsm"} and not any(k in ad for k in MEVZUAT_ADLARI):
            try:
                tespitler += kagit_oku(yol)
            except (ValueError, OSError) as h:
                atlanan.append(f"{yol.name} ({h})")
            continue
        if yol.suffix.lower() not in belge.DESTEKLENEN:
            atlanan.append(yol.name)
            continue
        try:
            metin = belge.metin_oku(yol)
        except belge.BelgeHatasi as h:
            atlanan.append(f"{yol.name} ({h})")
            continue
        kid = f"K{len(kaynaklar) + 1}"
        baslik, mlist = maddelere_bol(metin, kid)
        kaynaklar.append(f"{kid}: {baslik} ({yol.name})")
        maddeler += mlist
    return tespitler, maddeler, kaynaklar, atlanan


# ----------------------------------------------------------------------------
# Kontroller
# ----------------------------------------------------------------------------

def kokler(s: str) -> set[str]:
    return {w[:5] for w in katla(s).split() if len(w) >= 4 and not w.isdigit()}


def madde_sec(tespitler: list[Tespit], maddeler: list[Madde]) -> list[Madde]:
    if len(maddeler) <= TUM_MADDE_SINIRI:
        return maddeler
    secilen = {}
    km = {m.id: kokler(m.baslik + " " + m.metin) for m in maddeler}
    for t in tespitler:
        kt = kokler(t.konu + " " + t.metin)
        puan = sorted(((len(kt & km[m.id]), i) for i, m in enumerate(maddeler)), reverse=True)
        for p, i in puan[:TESPIT_BASINA_MADDE]:
            if p >= 2:
                secilen[maddeler[i].id] = maddeler[i]
    return [m for m in maddeler if m.id in secilen]


def kontrol_et(tespitler: list[Tespit], maddeler: list[Madde]) -> list[tuple[str, str, str]]:
    bulgular = []
    if not maddeler:
        bulgular.append(("orta", "Mevzuat / iç düzenleme metni verilmedi: kriterler (dayanak) müfettişçe eklenmeli", ""))
    sayi = Counter(t.no for t in tespitler)
    for no, n in sayi.items():
        if n > 1:
            bulgular.append(("yüksek", f"Tespit numarası tekrar ediyor: {no} ({n} kez)", no))
    for t in tespitler:
        if not t.kanit:
            bulgular.append(("orta", f"{t.no}: kanıt / çalışma kâğıdı referansı yok", t.no))
        if t.orneklem is not None and t.hatali is not None and t.hatali > t.orneklem:
            bulgular.append(("yüksek", f"{t.no}: hatalı adet ({t.hatali:g}) örneklemden ({t.orneklem:g}) büyük", t.no))
        if t.hatali is not None and t.orneklem is None:
            bulgular.append(("bilgi", f"{t.no}: örneklem büyüklüğü yazılmamış; hata oranı hesaplanamadı", t.no))
    grup = defaultdict(list)
    for t in tespitler:
        if t.konu:
            grup[(katla(t.birim), katla(t.konu))].append(t.no)
    for nolar in grup.values():
        if len(nolar) > 1:
            bulgular.append(("bilgi", f"Aynı birim ve konuda birden çok tespit: {', '.join(nolar)} — tek bulguda birleştirilebilir", ", ".join(nolar)))
    return bulgular


# ----------------------------------------------------------------------------
# Maskeleme
# ----------------------------------------------------------------------------

def takma_adlar(tespitler: list[Tespit], terimler: list[str]) -> dict[str, str]:
    harita = {}
    for i, ad in enumerate(dict.fromkeys(a for t in tespitler for a in t.personel), 1):
        harita[ad] = f"[PERSONEL-{i}]"
    for i, ad in enumerate(dict.fromkeys(a for t in tespitler for a in t.musteri), 1):
        harita.setdefault(ad, f"[MÜŞTERİ-{i}]")
    for i, ad in enumerate(dict.fromkeys(a.strip() for a in terimler if a.strip()), 1):
        harita.setdefault(ad, f"[KİŞİ-{i}]")
    return harita


def maskele(metin: str, harita: dict[str, str]) -> str:
    for gercek, takma in sorted(harita.items(), key=lambda x: -len(x[0])):
        metin = re.sub(re.escape(gercek), takma, metin, flags=re.I)
    return llm.maskele(metin)


def geri_ac(metin: str, harita: dict[str, str]) -> str:
    for gercek, takma in sorted(harita.items(), key=lambda x: -len(x[1])):
        metin = metin.replace(takma, gercek)
    return metin


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

SEMA = {
    "type": "object",
    "properties": {
        "bulgular": {"type": "array", "items": {
            "type": "object",
            "properties": {
                "baslik": {"type": "string"},
                "tespit_nolari": {"type": "array", "items": {"type": "string"}},
                "durum": {"type": "string"},
                "kriterler": {"type": "array", "items": {
                    "type": "object", "properties": {"madde": {"type": "string"}, "alinti": {"type": "string"}},
                    "required": ["madde", "alinti"], "additionalProperties": False}},
                "neden": {"type": "string"},
                "etki": {"type": "string"},
                "risk_duzeyi": {"type": "string", "enum": ["Yüksek", "Orta", "Düşük"]},
                "risk_gerekcesi": {"type": "string"},
                "oneriler": {"type": "array", "items": {"type": "string"}},
            },
            "required": ["baslik", "tespit_nolari", "durum", "kriterler", "neden", "etki", "risk_duzeyi", "risk_gerekcesi", "oneriler"],
            "additionalProperties": False}},
        "genel_not": {"type": "string"},
    },
    "required": ["bulgular", "genel_not"],
    "additionalProperties": False,
}


def duz(s: str) -> str:
    return " ".join(katla(s).split())


def sayilar(metin: str) -> set[str]:
    return {re.sub(r"\D", "", m) for m in re.findall(r"\d[\d.,]*\d", metin) if len(re.sub(r"\D", "", m)) >= 3}


def tespit_satiri(t: Tespit) -> str:
    parca = [f"[{t.no}] Birim: {t.birim or '—'} · Konu: {t.konu or '—'}", f"Tespit: {t.metin}"]
    if t.orneklem is not None or t.hatali is not None:
        oran = f" (hata oranı %{t.hata_orani * 100:.1f})".replace(".", ",") if t.hata_orani is not None else ""
        parca.append(f"Örneklem: {t.orneklem:g} · Hatalı: {t.hatali:g}{oran}" if t.orneklem is not None and t.hatali is not None
                     else f"Örneklem: {t.orneklem or '—'} · Hatalı: {t.hatali or '—'}")
    if t.tutar is not None:
        parca.append(f"Tutar: {tl(t.tutar)} TL")
    for et, d in (("Kanıt", t.kanit), ("Personel", ", ".join(t.personel)), ("Müşteri", ", ".join(t.musteri)),
                  ("Kök neden (müfettiş notu)", t.neden), ("Birim açıklaması", t.birim_aciklamasi)):
        if d:
            parca.append(f"{et}: {d}")
    return "\n".join(parca)


def hazirla(tespitler: list[Tespit], maddeler: list[Madde], kontroller, harita) -> tuple[str, str]:
    mesaj = "\n".join([
        "<tespitler>", *[maskele(tespit_satiri(t), harita) + "\n" for t in tespitler], "</tespitler>",
        "<kriter_metinleri>",
        *[f'<madde id="{m.id}" kaynak="{m.kaynak}" no="{m.no}" baslik="{m.baslik}">\n{m.metin}\n</madde>' for m in maddeler],
        "</kriter_metinleri>",
        "<kod_kontrolleri>", *([f"- [{o}] {maskele(a, harita)}" for o, a, _ in kontroller] or ["- Kontrol bulgusu yok"]), "</kod_kontrolleri>",
    ])
    return (BURASI / "prompt.md").read_text(encoding="utf-8"), mesaj


def denetle(y: dict, tespitler: list[Tespit], maddeler: list[Madde], mesaj: str) -> list[tuple[str, str, str]]:
    sorunlar = []
    mad = {m.id: m for m in maddeler}
    nolar = {t.no for t in tespitler}
    tmap = {t.no: t for t in tespitler}
    for k, b in enumerate(y["bulgular"], 1):
        b["sira"] = k
        uyd = [n for n in b["tespit_nolari"] if n not in nolar]
        if uyd:
            sorunlar.append(("yüksek", f"Bulgu {k}: tespitlerde olmayan numara yok sayıldı: {', '.join(uyd)}", f"Bulgu {k}"))
        b["tespit_nolari"] = [n for n in b["tespit_nolari"] if n in nolar]
        for n in b["tespit_nolari"]:
            tmap[n].bulgular.append(k)
        for kr in b["kriterler"]:
            m = mad.get(kr["madde"])
            if m is None:
                kr["durum"] = "Madde bulunamadı"
                sorunlar.append(("yüksek", f"Bulgu {k}: '{kr['madde']}' verilen metinlerde yok (uydurma dayanak)", f"Bulgu {k}"))
            elif not kr["alinti"].strip() or duz(kr["alinti"]) not in duz(m.metin):
                kr["durum"] = "Alıntı doğrulanamadı"
                sorunlar.append(("yüksek", f"Bulgu {k}: {kr['madde']} alıntısı maddede birebir geçmiyor", f"Bulgu {k}"))
            else:
                kr["durum"] = "Doğrulandı"
            kr["gosterim"] = f"{m.kaynak} md. {m.no}" if m and m.no else (m.kaynak if m else kr["madde"])
        b["dayanak_var"] = any(kr["durum"] == "Doğrulandı" for kr in b["kriterler"])
        if not b["dayanak_var"]:
            sorunlar.append(("orta", f"Bulgu {k} ({b['baslik']}): doğrulanmış dayanak yok — müfettişçe eklenmeli", f"Bulgu {k}"))
    for t in tespitler:
        if not t.bulgular:
            sorunlar.append(("orta", f"{t.no} hiçbir bulguda kullanılmamış", t.no))
    bilinen = sayilar(mesaj) | {x for x in re.findall(r"\d+", mesaj) if len(x) >= 3}
    metin = " ".join(" ".join([b["durum"], b["neden"], b["etki"], b["risk_gerekcesi"], *b["oneriler"]]) for b in y["bulgular"])
    y["dogrulanamayan_sayilar"] = sorted(s for s in sayilar(metin) if s not in bilinen)
    return sorunlar


def model_yaz(tespitler, maddeler, kontroller, harita) -> dict:
    sistem, mesaj = hazirla(tespitler, maddeler, kontroller, harita)
    y = llm.json_iste(sistem, mesaj, SEMA)
    y["sorunlar"] = denetle(y, tespitler, maddeler, mesaj)
    return y


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
ONEM = {"yüksek": "FDE2E1", "orta": "FFF4CE", "bilgi": "E8F0FE", "Yüksek": "FDE2E1", "Orta": "FFF4CE", "Düşük": "E8F0FE"}
KARAR = PatternFill("solid", fgColor="FFF4CE")
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def rapor_yaz(cikti: Path, tespitler, maddeler, kaynaklar, atlanan, kontroller, y, harita) -> Path:
    g = lambda s: geri_ac(s, harita)  # noqa: E731
    tum = kontroller + y["sorunlar"]
    wb = Workbook()
    o = wb.active
    o.title = "Özet"
    o.column_dimensions["A"].width, o.column_dimensions["B"].width = 30, 100
    risk = Counter(b["risk_duzeyi"] for b in y["bulgular"])
    for a, d in [("Tespit sayısı", len(tespitler)), ("Bulgu taslağı", len(y["bulgular"])),
                 ("Risk düzeyi (taslak)", f"Yüksek {risk['Yüksek']} · Orta {risk['Orta']} · Düşük {risk['Düşük']}"),
                 ("Dayanağı doğrulanmış bulgu", f"{sum(b['dayanak_var'] for b in y['bulgular'])} / {len(y['bulgular'])}"),
                 ("Kriter kaynakları", "\n".join(kaynaklar) or "verilmedi"), ("Modele gönderilen madde", len(maddeler)),
                 ("Okunamayan", "; ".join(atlanan) or "—"), ("Genel not (model)", g(y["genel_not"])),
                 ("Doğrulanamayan sayılar", ", ".join(y["dogrulanamayan_sayilar"]) or "—"), ("Model", llm.kullanim_ozeti())]:
        o.append([a, d])
        o.cell(o.max_row, 1).font = Font(bold=True)
        if a == "Doğrulanamayan sayılar" and y["dogrulanamayan_sayilar"]:
            o.cell(o.max_row, 2).fill = PatternFill("solid", fgColor=ONEM["yüksek"])
        for h in o[o.max_row]:
            h.alignment = UST

    bl = wb.create_sheet("Bulgular")
    _baslik(bl, ["No", "Başlık", "Tespitler", "Risk", "Durum (Tespit)", "Kriter (Dayanak)", "Neden", "Etki", "Risk Gerekçesi", "Öneriler",
                 "Müfettiş Onayı", "Birim Cevabı"], (5, 30, 10, 8, 50, 50, 36, 36, 30, 44, 18, 30))
    for b in y["bulgular"]:
        kriter = "\n".join(f"{kr['gosterim']}: \"{kr['alinti']}\"" + ("" if kr["durum"] == "Doğrulandı" else f" [{kr['durum']}]") for kr in b["kriterler"])
        bl.append([b["sira"], g(b["baslik"]), ", ".join(b["tespit_nolari"]), b["risk_duzeyi"], g(b["durum"]),
                   kriter if b["dayanak_var"] else (kriter + "\n" if kriter else "") + "DAYANAK MÜFETTİŞÇE EKLENMELİ", g(b["neden"]), g(b["etki"]),
                   g(b["risk_gerekcesi"]), "\n".join(f"- {g(x)}" for x in b["oneriler"]), "", ""])
        bl.cell(bl.max_row, 4).fill = PatternFill("solid", fgColor=ONEM[b["risk_duzeyi"]])
        if not b["dayanak_var"]:
            bl.cell(bl.max_row, 6).fill = PatternFill("solid", fgColor=ONEM["yüksek"])
        for c in (11, 12):
            bl.cell(bl.max_row, c).fill = KARAR
        for h in bl[bl.max_row]:
            h.alignment = UST

    ts = wb.create_sheet("Tespitler")
    _baslik(ts, ["Tespit No", "Birim", "Konu", "Tespit", "Örneklem", "Hatalı", "Hata Oranı", "Tutar", "Kanıt", "Bulgu"], (9, 14, 18, 60, 9, 8, 9, 14, 26, 8))
    for t in tespitler:
        ts.append([t.no, t.birim, t.konu, t.metin, t.orneklem and float(t.orneklem), t.hatali and float(t.hatali),
                   float(t.hata_orani) if t.hata_orani is not None else None, t.tutar and float(t.tutar), t.kanit, ", ".join(map(str, t.bulgular)) or "—"])
        ts.cell(ts.max_row, 7).number_format = "0.0%"
        ts.cell(ts.max_row, 8).number_format = "#,##0.00"
        ts.cell(ts.max_row, 4).alignment = UST
        if not t.bulgular:
            ts.cell(ts.max_row, 10).fill = PatternFill("solid", fgColor=ONEM["orta"])

    dk = wb.create_sheet("Dayanak Kontrolü")
    _baslik(dk, ["Bulgu", "Madde", "Alıntı", "Durum"], (7, 40, 90, 20))
    for b in y["bulgular"]:
        for kr in b["kriterler"]:
            dk.append([b["sira"], kr["gosterim"], kr["alinti"], kr["durum"]])
            dk.cell(dk.max_row, 4).fill = PatternFill("solid", fgColor="E3F4E1" if kr["durum"] == "Doğrulandı" else ONEM["yüksek"])
            dk.cell(dk.max_row, 3).alignment = UST

    kt = wb.create_sheet("Kontroller")
    _baslik(kt, ["Önem", "Kontrol", "İlgili", "İnceleme"], (9, 90, 16, 26))
    for on, a, il in tum:
        kt.append([on, a, il, ""])
        kt.cell(kt.max_row, 1).fill = PatternFill("solid", fgColor=ONEM[on])
        kt.cell(kt.max_row, 4).fill = KARAR
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

    md = cikti.with_suffix(".md")
    sat = ["# Teftiş Raporu Bulgu Taslakları", "", f"{len(tespitler)} tespit · {len(y['bulgular'])} bulgu taslağı · kriter kaynakları: " +
           ("; ".join(kaynaklar) or "verilmedi"), ""]
    for b in y["bulgular"]:
        sat += [f"## Bulgu {b['sira']}: {g(b['baslik'])}", "", f"**Risk düzeyi (taslak):** {b['risk_duzeyi']} — {g(b['risk_gerekcesi'])}  ",
                f"**İlgili tespitler:** {', '.join(b['tespit_nolari'])}", "", "**Tespit:**", "", g(b["durum"]), "", "**Kriter (dayanak):**", ""]
        sat += [f"- {kr['gosterim']}: \"{kr['alinti']}\"" + ("" if kr["durum"] == "Doğrulandı" else f" **[{kr['durum']}]**") for kr in b["kriterler"]]
        if not b["dayanak_var"]:
            sat += ["- **Dayanak müfettişçe eklenmeli.**"]
        sat += ["", f"**Neden:** {g(b['neden'])}", "", f"**Etki / risk:** {g(b['etki'])}", "", "**Öneriler:**", ""]
        sat += [f"- {g(x)}" for x in b["oneriler"]]
        sat += ["", "**Birim cevabı:** …", ""]
    if tum:
        sat += ["## Kontroller", ""] + [f"- **{o_}** {a}" for o_, a, _ in tum] + [""]
    if y["dogrulanamayan_sayilar"]:
        sat += [f"> Girdilerde bulunmayan sayılar (kontrol edin): {', '.join(y['dogrulanamayan_sayilar'])}", ""]
    sat += ["> Taslaktır; bulgu metinleri, risk düzeyleri ve dayanaklar müfettiş tarafından gözden geçirilmelidir."]
    md.write_text("\n".join(sat) + "\n", encoding="utf-8")
    return md


def calistir(klasor: Path, cikti: Path, terimler: list[str] | None = None, evet: bool = False) -> dict:
    if not klasor.is_dir():
        raise llm.LLMHatasi(f"{klasor} klasörü bulunamadı.")
    tespitler, maddeler, kaynaklar, atlanan = dosyalari_oku(klasor)
    if not tespitler:
        raise llm.LLMHatasi(f"{klasor}: çalışma kâğıdında tespit bulunamadı ('Tespit' sütunlu .csv/.xlsx).")
    kontroller = kontrol_et(tespitler, maddeler)
    secilen = madde_sec(tespitler, maddeler)
    harita = takma_adlar(tespitler, terimler or [])
    print(f"[OK] {len(tespitler)} tespit · {len(maddeler)} madde ({len(kaynaklar)} kaynak), {len(secilen)} madde gönderilecek")
    llm.onay_al(f"{len(tespitler)} tespit ve {len(secilen)} mevzuat maddesi gönderilecek "
                f"({sum(t.startswith('[PERSONEL') for t in harita.values())} personel, {sum(t.startswith('[MÜŞTERİ') for t in harita.values())} müşteri "
                "takma adlı; telefon, e-posta, TCKN, IBAN maskeli).", evet)
    y = model_yaz(tespitler, secilen, kontroller, harita)
    md = rapor_yaz(cikti, tespitler, secilen, kaynaklar, atlanan, kontroller, y, harita)
    return {"tespitler": tespitler, "maddeler": maddeler, "secilen": secilen, "kontroller": kontroller, "yanit": y, "harita": harita, "md": md}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Çalışma kâğıdındaki tespitlerden dayanak, risk ve öneri içeren teftiş bulgu taslakları hazırlar.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "ornek_sube",
                   help="Klasör: çalışma kâğıdı (.csv/.xlsx) ve mevzuat / iç düzenleme metinleri")
    p.add_argument("--gizle", nargs="*", default=[], metavar="AD", help="Ayrıca maskelenecek adlar")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "bulgu_taslaklari.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.cikti, a.gizle, a.evet)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    for o, ac, _ in s["kontroller"] + s["yanit"]["sorunlar"]:
        if o != "bilgi":
            print(f"[{'X' if o == 'yüksek' else '!'}] {ac}")
    if s["yanit"]["dogrulanamayan_sayilar"]:
        print(f"[!] Girdilerde olmayan sayılar: {', '.join(s['yanit']['dogrulanamayan_sayilar'])}")
    print(f"[OK] {len(s['yanit']['bulgular'])} bulgu taslağı: {a.cikti.resolve()} · {s['md'].resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
