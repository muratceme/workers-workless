"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import json
import os
import sys
import tempfile
import unittest
from datetime import date, datetime
from decimal import Decimal as D
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
SIP = ORNEK / "siparisler"
URUN = ORNEK / "urunler.csv"
BUGUN = date(2026, 10, 16)
CEVAP = json.loads((Path(__file__).parent / "sahte_cevaplar.json").read_text(encoding="utf-8"))
MESAJLAR = []


def kurgu_hatali() -> dict:
    """Sık model hataları: tutarsız toplamı 'düzeltmek', rengi çevirmek, olmayan ürün kodu uydurmak."""
    c = copy.deepcopy(CEVAP["kurgu"])
    c["satirlar"][0]["miktar"] = "1500"
    c["satirlar"][1]["renk"] = "Lacivert"
    c["satirlar"][2]["urun_kodu"] = "SW-200"
    return c


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    if "SP-88213" in kullanici:
        return copy.deepcopy(CEVAP["deneme"])
    return copy.deepcopy(CEVAP["kurgu"]) if "<duzeltme>" in kullanici else kurgu_hatali()


class KodTesti(unittest.TestCase):
    def test_yardimcilar(self):
        self.assertEqual((agent.tarih_coz("05.11.2026'ya kadar"), agent.tarih_coz("2026-11-20"), agent.tarih_coz("yakında")),
                         (date(2026, 11, 5), date(2026, 11, 20), None))
        self.assertEqual(agent.beden_sirasi(["XL", "S", "36", "XS", "M", "S"]), ["XS", "S", "M", "XL", "36"])
        self.assertEqual(agent.urunleri_oku(URUN)["FLN-100"], "FLN-100 Flanş")
        self.assertEqual(agent.urunleri_oku(None), {})

    def test_denetim(self):
        u = agent.urunleri_oku(URUN)
        m = llm.maskele((SIP / "kurgu_retail_po.txt").read_text(encoding="utf-8"))
        self.assertEqual(agent.denetle(CEVAP["kurgu"], m, u), [])
        s = agent.denetle(kurgu_hatali(), m, u)
        self.assertEqual(len(s), 3)
        for parca in ("Satır 1: Miktar '1500' belgede yazmıyor", "Satır 2: Renk 'Lacivert'", "Satır 3: ürün kodu 'SW-200'"):
            self.assertTrue(any(parca in x for x in s), parca)
        c = copy.deepcopy(CEVAP["kurgu"])
        c["satirlar"][0].update(teslim_tarihi="21.11.2026", bedenler=[{"beden": "XXL", "miktar": "300"}])
        c["satirlar"][1].update(miktar="", bedenler=[])
        s = agent.denetle(c, m, u)
        self.assertTrue(any("21.11.2026 belgede yazmıyor" in x for x in s) and any("beden 'XXL'" in x for x in s) and any("ikisi" in x or "boş" in x for x in s))
        self.assertIn("ürün listesi verilmedi", agent.denetle(CEVAP["kurgu"], m, {})[0])
        self.assertEqual(agent.denetle({"siparis_var": False}, m, u), [])

    def test_bulgular(self):
        u = agent.urunleri_oku(URUN)
        k = copy.deepcopy(CEVAP["kurgu"])
        agent.bulgular(k, u, BUGUN, 21)
        self.assertEqual([x["_durum"] for x in k["satirlar"]], ["Hatalı", "Uygun", "Kontrol"])
        self.assertEqual(k["satirlar"][0]["_bulgu"][0][1], "Beden toplamı 1.500, satır toplamı 1.600: belge tutarsız, müşteriye teyit ettirin")
        self.assertEqual(k["_bulgu"], [])                                     # 1.600 + 1.150 + 750 = 3.500 belgedeki genel toplam
        k = copy.deepcopy(CEVAP["kurgu"]) | {"genel_toplam": "3.600"}
        agent.bulgular(k, u, BUGUN, 21)
        self.assertEqual(k["_bulgu"], [("Hata", "Satırlar toplamı 3.500, belgedeki genel toplam 3.600")])
        d = copy.deepcopy(CEVAP["deneme"])
        agent.bulgular(d, u, BUGUN, 21)
        self.assertEqual([[o for o, _ in x["_bulgu"]] for x in d["satirlar"]], [["Dikkat"], ["Dikkat", "Bilgi"], ["Dikkat", "Dikkat", "Bilgi"]])
        d = copy.deepcopy(CEVAP["deneme"])
        agent.bulgular(d, u, date(2026, 11, 4), 0)
        self.assertEqual(d["satirlar"][2]["_bulgu"][1], ("Hata", "Teslim tarihi geçmiş (25.10.2026)"))
        k = copy.deepcopy(CEVAP["kurgu"])
        k["satirlar"][1]["miktar"] = ""                                       # satır toplamı yoksa beden toplamı kullanılır
        agent.bulgular(k, u, BUGUN, 21)
        self.assertEqual(k["satirlar"][1]["_miktar"], D(1150))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ornek_ve_duzeltme(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "s.xlsx"
            self.assertEqual(agent.main(["--evet", "--cikti", str(cikti)]), 0)
            self.assertEqual(len(MESAJLAR), 3)                          # 2 belge + Kurgu Retail için 1 düzeltme
            self.assertIn("TS-101 | Erkek basic tişört", MESAJLAR[0])
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Sipariş Satırları", "Üretim Siparişleri", "Beden Dağılımı", "T&A Siparişleri", "Kontrol"])
            u = [list(r) for r in wb["Üretim Siparişleri"].iter_rows(values_only=True)]
            self.assertEqual(u[0], ["Sipariş No", "Müşteri", "Ürün", "Miktar", "Termin", "Öncelik"])
            self.assertEqual(u[1][:5], ["SP-88213/1", "Deneme Makine San.", "FLN-100 Flanş", 400, datetime(2026, 11, 4)])
            b = [list(r) for r in wb["Beden Dağılımı"].iter_rows(values_only=True)]
            self.assertEqual(b[0], ["Sipariş No", "Model", "Renk", "XS", "S", "M", "L", "XL", "Toplam"])
            self.assertEqual(b[1], ["KR-4471", "TS-101", "White", None, 300, 500, 500, 200, 1500])
            t = [list(r) for r in wb["T&A Siparişleri"].iter_rows(values_only=True)]
            self.assertIn(["KR-4471/TS-101", "Kurgu Retail GmbH", "TS-101", 2750, datetime(2026, 10, 12), datetime(2026, 11, 20), None], t)
            self.assertEqual(wb["Sipariş Satırları"].cell(5, 17).value, "Hatalı")

    @mock.patch.object(llm, "json_iste", side_effect=lambda *a, **k: {"siparis_var": False, "musteri": "", "po_no": "", "po_tarihi": "",
                                                                     "para_birimi": "Belirtilmemiş", "teslim_sekli": "", "satirlar": [], "genel_toplam": "",
                                                                     "notlar": ["Bu bir teklif istemi."]})
    def test_siparis_degil_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "t.txt"
            y.write_text("Fiyat teklifi rica ederiz: 500 adet flanş.\n", encoding="utf-8")
            s = agent.calistir(y, Path(tmp) / "s.xlsx", None, BUGUN, evet=True)
            self.assertFalse(s["belgeler"][0]["c"]["siparis_var"])
            self.assertNotIn("Beden Dağılımı", load_workbook(Path(tmp) / "s.xlsx").sheetnames)
            self.assertEqual(agent.main(["--min-gun", "-1", "--evet"]), 2)
            self.assertEqual(agent.main(["--girdi", str(Path(tmp) / "yok"), "--evet"]), 1)
            self.assertEqual(agent.main(["--urunler", str(Path(tmp) / "yok.csv"), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
