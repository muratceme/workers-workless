"""
Müşteri Siparişi Bilgi Çıkarma — Gerisi Bende AI Agent
Üretim Planlama › Üretim Planlama Uzmanı (tekstil Planlama Uzmanı da kullanır)

Müşterilerden e-posta / PDF ile gelen siparişleri (PO) planlama tablolarına aktarır.
1. Kod sipariş belgelerini ve (varsa) ürün listenizi okur; telefon, e-posta, IBAN ve kimlik numaralarını maskeler.
2. Model her siparişin müşterisini, PO no ve tarihini, para birimini, teslim şeklini ve satırlarını çıkarır: müşterinin
   ürün kodu / açıklaması, sizin ürün kodunuz (ürün listesinden eşleştirme), renk, beden dağılımı, miktar, birim fiyat,
   teslim tarihi. Belgede yazan satır ve genel toplamları da aktarır.
3. Kod çıkarımı denetler: metinler ve sayılar belgede geçiyor mu, ürün kodu listede var mı, teslim tarihi okunuyor mu.
   Sorun varsa modele geri gönderilir (en fazla --tur kez).
4. Kod bulguları hesaplar: beden toplamı ≠ satır toplamı, satırlar toplamı ≠ belgedeki genel toplam, ürün kartı yok,
   teslim tarihi geçmiş ya da --min-gun içinde, fiyat yok.
Çıktı: "Üretim Siparişleri" (Haftalık Üretim Çizelgesi girdisi), "Beden Dağılımı" (Kumaş ve Aksesuar İhtiyaç Hesabı
girdisi) ve "T&A Siparişleri" (Sipariş Termin Planı girdisi) sayfaları.

Kullanım:
    python agent.py                                                    # örnek: tekstil PO ve metal parça siparişi
    python agent.py --girdi ./siparisler --urunler urun_listesi.xlsx --min-gun 21
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
ORNEK_TARIH = date(2026, 10, 16)
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
RENK = {"Hata": "FDE2E1", "Dikkat": "FFF4CE", "Bilgi": "E8F0FE", "Hatalı": "FDE2E1", "Kontrol": "FFF4CE", "Uygun": "E3F5E1"}
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}
BEDEN_SIRA = ["XXS", "XS", "S", "M", "L", "XL", "XXL", "2XL", "3XL", "4XL"]


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1 or (s.count(".") == 1 and len(s.split(".")[1]) == 3):
        s = s.replace(".", "")
    try:
        return Decimal(s) if s not in ("", "-", ".") else None
    except InvalidOperation:
        return None


def sayilar(metin: str) -> set[Decimal]:
    return {v for x in re.findall(r"\d+(?:[.,]\d+)*", str(metin)) if (v := para(x)) is not None}


def tarih_coz(x) -> date | None:
    s = str(x or "").strip()
    m = re.search(r"\b(\d{1,2})[./-](\d{1,2})[./-](\d{4})\b", s)
    if m:
        try:
            return date(int(m[3]), int(m[2]), int(m[1]))
        except ValueError:
            return None
    m = re.search(r"\b(\d{4})-(\d{2})-(\d{2})\b", s)
    return date(int(m[1]), int(m[2]), int(m[3])) if m else None


def tr(x: Decimal) -> str:
    s = f"{x:,.2f}".rstrip("0").rstrip(".")
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def urunleri_oku(yol: Path | None) -> dict[str, str]:
    if yol is None:
        return {}
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: ürün listesi bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, read_only=True, data_only=True)
        s = [[str(c or "").strip() for c in r] for r in wb.active.iter_rows(values_only=True)]
        wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=metin.split("\n")[0].count)))
    if not s or len(s[0]) < 2:
        raise llm.LLMHatasi(f"{yol.name}: 'Ürün Kodu' ve 'Ürün Adı' sütunları gerekli.")
    return {r[0].strip(): r[1].strip() for r in s[1:] if len(r) >= 2 and r[0].strip()}


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

BEDEN = {"type": "object", "properties": {"beden": S, "miktar": S}, "required": ["beden", "miktar"], "additionalProperties": False}
SATIR = {"type": "object", "properties": {"musteri_kodu": S, "aciklama": S, "urun_kodu": S, "renk": S, "bedenler": {"type": "array", "items": BEDEN},
                                          "miktar": S, "birim": S, "birim_fiyat": S, "teslim_tarihi": S},
         "required": ["musteri_kodu", "aciklama", "urun_kodu", "renk", "bedenler", "miktar", "birim", "birim_fiyat", "teslim_tarihi"], "additionalProperties": False}
SEMA = {"type": "object", "properties": {
    "siparis_var": {"type": "boolean"}, "musteri": S, "po_no": S, "po_tarihi": S, "para_birimi": {"type": "string", "enum": ["TL", "USD", "EUR", "GBP", "Diğer", "Belirtilmemiş"]},
    "teslim_sekli": S, "satirlar": {"type": "array", "items": SATIR}, "genel_toplam": S, "notlar": {"type": "array", "items": S}},
    "required": ["siparis_var", "musteri", "po_no", "po_tarihi", "para_birimi", "teslim_sekli", "satirlar", "genel_toplam", "notlar"], "additionalProperties": False}


def denetle(c: dict, metin: str, urunler: dict[str, str]) -> list[str]:
    if not c["siparis_var"]:
        return []
    sorun, say = [], sayilar(metin)

    def sayi(on, ad, v):
        if v.strip() and (para(v) is None or para(v) not in say):
            sorun.append(f"{on}: {ad} '{v}' belgede yazmıyor (hesaplamayın; belgede yoksa boş bırakın)")

    for a, ad in (("musteri", "Müşteri"), ("po_no", "PO no"), ("po_tarihi", "PO tarihi"), ("teslim_sekli", "Teslim şekli")):
        if c[a].strip() and not metinde_var(c[a], metin):
            sorun.append(f"Sipariş: {ad} '{c[a][:50]}' belgede geçmiyor")
    sayi("Sipariş", "Genel toplam", c["genel_toplam"])
    if not c["satirlar"]:
        sorun.append("Sipariş: satır yok")
    for j, s in enumerate(c["satirlar"], 1):
        on = f"Satır {j}"
        for a, ad in (("musteri_kodu", "Müşteri kodu"), ("aciklama", "Açıklama"), ("renk", "Renk")):
            if s[a].strip() and not metinde_var(s[a], metin):
                sorun.append(f"{on}: {ad} '{s[a][:50]}' belgede geçmiyor")
        for a, ad in (("miktar", "Miktar"), ("birim_fiyat", "Birim fiyat")):
            sayi(on, ad, s[a])
        for b in s["bedenler"]:
            if not metinde_var(b["beden"], metin):
                sorun.append(f"{on}: beden '{b['beden']}' belgede geçmiyor")
            sayi(on, f"{b['beden']} bedeni", b["miktar"])
        if not s["miktar"].strip() and not s["bedenler"]:
            sorun.append(f"{on}: miktar da beden dağılımı da boş")
        if s["teslim_tarihi"].strip():
            t = tarih_coz(s["teslim_tarihi"])
            if t is None:
                sorun.append(f"{on}: teslim tarihi '{s['teslim_tarihi']}' GG.AA.YYYY olarak okunamadı")
            elif f"{t:%d.%m.%Y}" not in metin and t.isoformat() not in metin:
                sorun.append(f"{on}: teslim tarihi {t:%d.%m.%Y} belgede yazmıyor")
        if urunler and s["urun_kodu"].strip() and s["urun_kodu"] not in urunler:
            sorun.append(f"{on}: ürün kodu '{s['urun_kodu']}' ürün listesinde yok; eşleşme yoksa boş bırakın")
        if not urunler and s["urun_kodu"].strip():
            sorun.append(f"{on}: ürün listesi verilmedi; urun_kodu boş olmalı")
    return sorun


def bulgular(c: dict, urunler: dict[str, str], bugun: date, min_gun: int) -> None:
    """Kod: satır ve sipariş bulguları (yerinde ekler)."""
    c["_bulgu"] = []
    toplam = Decimal(0)
    for s in c["satirlar"]:
        b = []
        m = para(s["miktar"])
        bt = sum((para(x["miktar"]) or Decimal(0) for x in s["bedenler"]), Decimal(0))
        if s["bedenler"] and m is not None and bt != m:
            b.append(("Hata", f"Beden toplamı {tr(bt)}, satır toplamı {tr(m)}: belge tutarsız, müşteriye teyit ettirin"))
        s["_miktar"] = m if m is not None else (bt if s["bedenler"] else None)
        toplam += s["_miktar"] or Decimal(0)
        if urunler and not s["urun_kodu"].strip():
            b.append(("Dikkat", "Ürün listesinde karşılığı yok (yeni ürün olabilir): ürün kartı / reçete açılmalı"))
        t = tarih_coz(s["teslim_tarihi"])
        if t is None:
            b.append(("Dikkat", "Teslim tarihi yok"))
        elif t < bugun:
            b.append(("Hata", f"Teslim tarihi geçmiş ({t:%d.%m.%Y})"))
        elif (t - bugun).days < min_gun:
            b.append(("Dikkat", f"Teslime {(t - bugun).days} gün var (en az {min_gun} gün istenir): kapasite ve malzeme teyidi gerekir"))
        if not s["birim_fiyat"].strip():
            b.append(("Bilgi", "Birim fiyat siparişte yok: teklif / fiyat listesiyle karşılaştırın"))
        s["_teslim"] = t
        s["_bulgu"] = b
        onem = {o for o, _ in b}
        s["_durum"] = "Hatalı" if "Hata" in onem else "Kontrol" if "Dikkat" in onem else "Uygun"
    g = para(c["genel_toplam"])
    if g is not None and g != toplam:
        c["_bulgu"].append(("Hata", f"Satırlar toplamı {tr(toplam)}, belgedeki genel toplam {tr(g)}"))
    if c["para_birimi"] == "Belirtilmemiş" and any(s["birim_fiyat"].strip() for s in c["satirlar"]):
        c["_bulgu"].append(("Dikkat", "Fiyat var ama para birimi belirtilmemiş"))


def uret(metin: str, urunler: dict[str, str], duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    liste = "\n".join(f"{k} | {v}" for k, v in urunler.items()) or "(verilmedi: urun_kodu alanlarını boş bırakın)"
    mesaj = f"<urun_listesi>\n{liste}\n</urun_listesi>\n<siparis>\n{metin}\n</siparis>"
    if duzeltme:
        mesaj += ("\n<duzeltme>\nÖnceki çıkarım: " + json.dumps(onceki, ensure_ascii=False) + "\nKodun bulduğu sorunlar:\n"
                  + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")
    return llm.json_iste(sistem, mesaj, SEMA)


def duzeltme_dongusu(uret_f, denetle_f, tur: int) -> tuple[dict, list[str], int]:
    taslak = uret_f(None, None)
    kalan = denetle_f(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret_f(kalan, taslak)
        kalan = denetle_f(taslak)
    return taslak, kalan, yapilan


# ----------------------------------------------------------------------------
# Çıktı
# ----------------------------------------------------------------------------

def _baslik(ws, basliklar, genislik=()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font, h.alignment = BASLIK_DOLGU, BASLIK_YAZI, UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def _f(x):
    return None if x is None else float(x)


def _model(s: dict, urunler: dict[str, str]) -> str:
    return s["urun_kodu"] or s["musteri_kodu"] or s["aciklama"][:30]


def beden_sirasi(bedenler: list[str]) -> list[str]:
    return sorted(dict.fromkeys(bedenler), key=lambda b: (BEDEN_SIRA.index(b.upper()) if b.upper() in BEDEN_SIRA else 100, b))


def rapor_yaz(cikti: Path, belgeler: list[dict], urunler: dict[str, str]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sipariş Satırları"
    _baslik(ws, ["Kaynak Dosya", "Müşteri", "PO No", "PO Tarihi", "Ürün Kodu", "Ürün Adı (bizim)", "Müşteri Kodu", "Açıklama (müşteri)", "Renk", "Beden Dağılımı",
                 "Miktar", "Birim", "Birim Fiyat", "Para Birimi", "Teslim Tarihi", "Teslim Şekli", "Durum", "Bulgular"],
            (26, 24, 14, 11, 10, 30, 12, 40, 12, 30, 9, 7, 10, 7, 11, 14, 9, 60))
    uretim, beden_satir, tna = [], [], []
    for d in belgeler:
        c = d["c"]
        if not c["siparis_var"]:
            continue
        for sira, s in enumerate(c["satirlar"], 1):
            ws.append([d["dosya"], c["musteri"], c["po_no"], tarih_coz(c["po_tarihi"]), s["urun_kodu"], urunler.get(s["urun_kodu"], ""), s["musteri_kodu"],
                       s["aciklama"], s["renk"], " · ".join(f"{b['beden']}: {b['miktar']}" for b in s["bedenler"]), _f(s["_miktar"]), s["birim"],
                       _f(para(s["birim_fiyat"])), c["para_birimi"], s["_teslim"], c["teslim_sekli"], s["_durum"],
                       "\n".join(f"{o}: {m}" for o, m in s["_bulgu"] + c["_bulgu"])])
            n = ws.max_row
            ws.cell(n, 4).number_format = ws.cell(n, 15).number_format = "DD.MM.YYYY"
            for j in (2, 3, 5, 7, 8, 9, 10, 11, 12, 13, 15, 16):
                ws.cell(n, j).fill = AI_DOLGU
            ws.cell(n, 17).fill = PatternFill("solid", fgColor=RENK[s["_durum"]])
            for h in ws[n]:
                h.alignment = UST
            ad = urunler.get(s["urun_kodu"]) or s["aciklama"]
            no = f"{c['po_no']}/{sira}" if len(c["satirlar"]) > 1 else c["po_no"]          # çizelgede sipariş no tekil olmalı
            uretim.append([no, c["musteri"], ad, _f(s["_miktar"]), s["_teslim"], ""])
            if s["bedenler"]:
                beden_satir.append((c["po_no"], _model(s, urunler), s["renk"], {b["beden"]: para(b["miktar"]) for b in s["bedenler"]}, s["_miktar"]))
        modeller: dict[str, list] = {}
        for s in c["satirlar"]:
            k = _model(s, urunler)
            m = modeller.setdefault(k, [Decimal(0), s["_teslim"]])
            m[0] += s["_miktar"] or Decimal(0)
            if s["_teslim"] and (m[1] is None or s["_teslim"] < m[1]):
                m[1] = s["_teslim"]
        for k, (adet, sevk) in modeller.items():
            tna.append([c["po_no"] if len(modeller) == 1 else f"{c['po_no']}/{k}", c["musteri"], k, _f(adet), tarih_coz(c["po_tarihi"]), sevk, None])

    u = wb.create_sheet("Üretim Siparişleri")
    _baslik(u, ["Sipariş No", "Müşteri", "Ürün", "Miktar", "Termin", "Öncelik"], (16, 24, 32, 10, 12, 8))
    for r in uretim:
        u.append(r)
        u.cell(u.max_row, 5).number_format = "DD.MM.YYYY"
    if beden_satir:
        bedenler = beden_sirasi([b for _, _, _, d, _ in beden_satir for b in d])
        bd = wb.create_sheet("Beden Dağılımı")
        _baslik(bd, ["Sipariş No", "Model", "Renk", *bedenler, "Toplam"])
        for po, model, renk, d, top in beden_satir:
            bd.append([po, model, renk, *(_f(d.get(b)) for b in bedenler), _f(sum((v or 0 for v in d.values()), Decimal(0)))])
    t = wb.create_sheet("T&A Siparişleri")
    _baslik(t, ["Sipariş No", "Müşteri", "Model", "Adet", "PO Tarihi", "Sevk Tarihi", "Kapasite Dikim"], (20, 24, 16, 9, 11, 11, 10))
    for r in tna:
        t.append(r)
        t.cell(t.max_row, 5).number_format = t.cell(t.max_row, 6).number_format = "DD.MM.YYYY"
    k = wb.create_sheet("Kontrol")
    _baslik(k, ["Dosya", "Müşteri", "PO No", "Satır", "Genel Toplam (belge)", "Sipariş Bulguları", "Model Notları", "Düzeltme Turu", "Son Kontrol"],
            (26, 24, 14, 6, 12, 50, 50, 8, 50))
    for d in belgeler:
        c = d["c"]
        k.append([d["dosya"], c.get("musteri", ""), c.get("po_no", ""), len(c.get("satirlar", [])), c.get("genel_toplam", ""),
                  "\n".join(f"{o}: {m}" for o, m in c.get("_bulgu", [])) if c["siparis_var"] else "Belgede sipariş bulunamadı",
                  "\n".join(c.get("notlar", [])), d["tur"], "; ".join(d["kalan"]) or "Geçti"])
        k.cell(k.max_row, 9).fill = KIRMIZI if d["kalan"] else YESIL
        for h in k[k.max_row]:
            h.alignment = UST
    k.append([])
    k.append([f"Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    k.append(["Önemli: Mor hücreler yapay zekâ ile çıkarılmıştır; siparişi onaylamadan önce PO'nun aslıyla karşılaştırın. Beden ve genel toplam "
              "kontrolleri, termin süresi ve ürün kartı eşleşmesi kod tarafından yapılmıştır. 'Üretim Siparişleri' Haftalık Üretim Çizelgesi'ne, "
              "'Beden Dağılımı' Kumaş ve Aksesuar İhtiyaç Hesabı'na, 'T&A Siparişleri' Sipariş Termin Planı'na girdi olarak verilebilir "
              "(gerekirse sayfayı ayrı dosyaya kopyalayın)."])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def dosyalar(girdi: Path) -> list[Path]:
    if girdi.is_dir():
        l = sorted(p for p in girdi.iterdir() if p.is_file() and p.suffix.lower() in belge.DESTEKLENEN)
        if not l:
            raise belge.BelgeHatasi(f"{girdi}: klasörde .pdf, .docx veya .txt dosyası yok.")
        return l
    if not girdi.exists():
        raise belge.BelgeHatasi(f"{girdi}: dosya bulunamadı.")
    return [girdi]


def calistir(girdi: Path, cikti: Path, urun_yolu: Path | None = None, bugun: date = ORNEK_TARIH, min_gun: int = 21, tur: int = 2,
             evet: bool = False) -> dict:
    urunler = urunleri_oku(urun_yolu)
    l = dosyalar(girdi)
    print(f"[OK] {len(l)} sipariş belgesi · ürün listesi: {len(urunler)}")
    llm.onay_al(f"{len(l)} sipariş belgesinin metni ve ürün listesi gönderilecek (telefon, e-posta, IBAN, TCKN maskeli).", evet)
    belgeler = []
    for y in l:
        ham = belge.metin_oku(y)
        if len(ham.strip()) < 20:
            raise belge.BelgeHatasi(f"{y.name}: metin bulunamadı.")
        metin = llm.maskele(ham)
        c, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(metin, urunler, d, o), lambda x: denetle(x, metin, urunler), tur)
        if c["siparis_var"]:
            bulgular(c, urunler, bugun, min_gun)
        belgeler.append({"dosya": y.name, "c": c, "kalan": kalan, "tur": yapilan})
    rapor_yaz(cikti, belgeler, urunler)
    return {"belgeler": belgeler}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Müşteri siparişlerini (PO) planlama tablolarına aktarır; beden, toplam ve termin kontrollerini kodla yapar.")
    p.add_argument("--girdi", type=Path, default=ORNEK / "siparisler", help="Sipariş belgesi ya da klasörü (.pdf, .docx, .txt)")
    p.add_argument("--urunler", type=Path, default=None, help="Ürün listesi (.csv/.xlsx: Ürün Kodu, Ürün Adı)")
    p.add_argument("--tarih", default=None, help="Bugünün tarihi GG.AA.YYYY (varsayılan bugün; örnekte 16.10.2026)")
    p.add_argument("--min-gun", type=int, default=21, help="Teslime bundan az gün varsa uyar (varsayılan 21)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "musteri_siparisleri.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    ornek = a.girdi == ORNEK / "siparisler"
    urunler = a.urunler or (ORNEK / "urunler.csv" if ornek else None)
    bugun = tarih_coz(a.tarih) if a.tarih else (ORNEK_TARIH if ornek else date.today())
    if bugun is None or a.min_gun < 0:
        print("[X] --tarih GG.AA.YYYY biçiminde, --min-gun 0 ya da büyük olmalı")
        return 2
    try:
        s = calistir(a.girdi, a.cikti, urunler, bugun, a.min_gun, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    for d in s["belgeler"]:
        c = d["c"]
        if not c["siparis_var"]:
            print(f"[!] {d['dosya']}: sipariş bulunamadı")
            continue
        print(f"[OK] {c['po_no']} ({c['musteri']}): {len(c['satirlar'])} satır")
        for o, m in c["_bulgu"] + [(o, f"{x['aciklama'][:30]}: {m}") for x in c["satirlar"] for o, m in x["_bulgu"]]:
            if o == "Hata":
                print(f"[!]   {m}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
