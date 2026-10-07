"""
Mülakat Değerlendirme Özeti — Workers / Workless AI Agent
İnsan Kaynakları › İşe Alım Uzmanı

1. Kod, mülakat puan ve notlarını okur: aday × yetkinlik ortalaması, yetkinlik ağırlıklarıyla toplam puan,
   sıralama, mülakatçılar arası puan farkı (≥ 2 puan → kalibrasyon) ve eksik puanlar.
2. Notlar ortak ayrımcılık tarayıcısından geçer; korunan özelliğe (yaş, medeni hâl, aile, sağlık, din, köken...)
   dayalı yorum içeren notlar "değerlendirme dışı" işaretlenir ve bu puanlar hariç tutularak ikinci bir
   toplam puan da hesaplanır.
3. Aday ve mülakatçı adları modele gitmeden takma adla (A1, M1) değiştirilir. Model her aday için kanıta
   dayalı güçlü yönler, gelişim alanları, görüş ayrılıkları, doğrulanacak noktalar ve öneri taslağı yazar.
   Rapor şirket içinde kaldığı için takma adlar raporda gerçek adlara geri çevrilir.

Kullanım:
    python agent.py                                              # örnek notlarla dener
    python agent.py --girdi mulakat_notlari.xlsx --yetkinlikler yetkinlikler.xlsx
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import OrderedDict, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import ayrimcilik
import llm

BURASI = Path(__file__).resolve().parent
UYUMSUZLUK = 2


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def sayi_yaz(x) -> str:
    return "-" if x is None else f"{x:.2f}"


def puan(x) -> float | None:
    try:
        v = float(str(x).replace(",", "."))
    except (TypeError, ValueError):
        return None
    return v if 1 <= v <= 5 else None


def notlari_oku(yol: Path) -> list[dict]:
    s = tablo_oku(yol)
    b = [kucuk(x) for x in s[0]]
    bul = lambda *a: next((i for i, x in enumerate(b) if x in a), None)  # noqa: E731
    k = {"aday": bul("aday", "aday adı"), "mulakatci": bul("mülakatçı", "mulakatci", "değerlendiren"),
         "yetkinlik": bul("yetkinlik", "kriter"), "puan": bul("puan", "skor"), "not": bul("not", "notlar", "yorum", "kanıt")}
    if None in (k["aday"], k["mulakatci"], k["yetkinlik"], k["puan"]):
        raise SystemExit(f"Aday, Mülakatçı, Yetkinlik ve Puan sütunları gerekli. Başlıklar: {s[0]}")
    al = lambda r, a: r[k[a]] if k[a] is not None and k[a] < len(r) else None  # noqa: E731
    return [{"aday": str(al(r, "aday")).strip(), "mulakatci": str(al(r, "mulakatci")).strip(), "yetkinlik": str(al(r, "yetkinlik")).strip(),
             "puan": puan(al(r, "puan")), "not": str(al(r, "not") or "").strip()} for r in s[1:] if al(r, "aday")]


def agirliklari_oku(yol: Path | None) -> dict[str, float]:
    if not yol:
        return {}
    s = tablo_oku(yol)
    b = [kucuk(x) for x in s[0]]
    i_ad = next((i for i, x in enumerate(b) if x in ("yetkinlik", "ad")), 0)
    i_ag = next((i for i, x in enumerate(b) if x in ("ağırlık", "ağırlık %", "önem")), None)
    if i_ag is None:
        return {}
    return {str(r[i_ad]).strip(): float(str(r[i_ag]).replace(",", ".")) for r in s[1:] if r[i_ad] and r[i_ag] not in (None, "")}


# ----------------------------------------------------------------------------
# Hesap (deterministik)
# ----------------------------------------------------------------------------

def hesapla(notlar: list[dict], agirlik: dict[str, float]) -> dict:
    for n in notlar:
        n["bulgular"] = [b for b in ayrimcilik.tara(n["not"]) if b["seviye"] == "yüksek"]
    adaylar = list(dict.fromkeys(n["aday"] for n in notlar))
    yetkinlikler = list(dict.fromkeys(list(agirlik) + [n["yetkinlik"] for n in notlar]))
    mulakatcilar = list(dict.fromkeys(n["mulakatci"] for n in notlar))
    sonuc = OrderedDict()
    for a in adaylar:
        satir = {"yetkinlik": OrderedDict(), "eksik": [], "uyumsuz": [], "onyargi": []}
        for y in yetkinlikler:
            ns = [n for n in notlar if n["aday"] == a and n["yetkinlik"] == y]
            puanlar = [n["puan"] for n in ns if n["puan"] is not None]
            temiz = [n["puan"] for n in ns if n["puan"] is not None and not n["bulgular"]]
            for m in mulakatcilar:
                if any(n["mulakatci"] == m for n in notlar if n["aday"] == a) and not any(n["mulakatci"] == m and n["puan"] is not None for n in ns):
                    satir["eksik"].append(f"{m} → {y}")
            fark = max(puanlar) - min(puanlar) if len(puanlar) > 1 else 0
            if fark >= UYUMSUZLUK:
                satir["uyumsuz"].append((y, fark, {n["mulakatci"]: n["puan"] for n in ns}))
            satir["onyargi"] += [(y, n["mulakatci"], b["ifade"]) for n in ns for b in n["bulgular"]]
            satir["yetkinlik"][y] = {"ort": sum(puanlar) / len(puanlar) if puanlar else None,
                                     "ort_temiz": sum(temiz) / len(temiz) if temiz else None, "fark": fark, "n": len(puanlar)}

        def toplam(alan):
            pay = payda = 0.0
            for y, v in satir["yetkinlik"].items():
                if v[alan] is not None:
                    w = agirlik.get(y, 1.0)
                    pay += w * v[alan]
                    payda += w
            return pay / payda if payda else None
        satir["toplam"] = toplam("ort")
        satir["toplam_temiz"] = toplam("ort_temiz")
        sonuc[a] = satir
    sira = sorted(sonuc, key=lambda a: -(sonuc[a]["toplam_temiz"] or 0))
    for i, a in enumerate(sira, 1):
        sonuc[a]["sira"] = i
    return {"adaylar": sonuc, "yetkinlikler": yetkinlikler, "mulakatcilar": mulakatcilar}


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

SEMA = {
    "type": "object",
    "properties": {"adaylar": {"type": "array", "items": {
        "type": "object",
        "properties": {
            "aday": {"type": "string"},
            "guclu_yonler": {"type": "array", "items": {"type": "string"}},
            "gelisim_alanlari": {"type": "array", "items": {"type": "string"}},
            "gorus_ayriliklari": {"type": "array", "items": {"type": "string"}},
            "dogrulanacaklar": {"type": "array", "items": {"type": "string"}},
            "oneri": {"type": "string", "enum": ["ilerlet", "beklet", "ilerletme"]},
            "oneri_gerekce": {"type": "string"},
        },
        "required": ["aday", "guclu_yonler", "gelisim_alanlari", "gorus_ayriliklari", "dogrulanacaklar", "oneri", "oneri_gerekce"],
        "additionalProperties": False}}},
    "required": ["adaylar"],
    "additionalProperties": False,
}


def takma(notlar, h) -> dict[str, str]:
    harita = {}
    for i, a in enumerate(dict.fromkeys(n["aday"] for n in notlar), 1):
        harita[a] = f"A{i}"
    for i, m in enumerate(h["mulakatcilar"], 1):
        harita[m] = f"M{i}"
    return harita


def gizle(metin: str, harita: dict[str, str]) -> str:
    metin = llm.maskele(metin)
    for gercek in sorted(harita, key=len, reverse=True):
        metin = re.sub(re.escape(gercek), harita[gercek], metin, flags=re.I)
    return metin


def ac(metin: str, harita: dict[str, str]) -> str:
    for gercek, t in sorted(harita.items(), key=lambda i: -len(i[1])):
        metin = re.sub(rf"\b{re.escape(t)}\b", gercek, metin)
    return metin


def ozetle(notlar: list[dict], h: dict, harita: dict[str, str]) -> dict[str, dict]:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    po = []
    for a, x in h["adaylar"].items():
        po.append(f"{harita[a]}: toplam {sayi_yaz(x['toplam'])}" + (f", değerlendirme dışı notlar hariç {sayi_yaz(x['toplam_temiz'])}" if x["onyargi"] else "")
                  + "; " + "; ".join(f"{y} ort {v['ort']:.2f}" + (f" (fark {v['fark']:g})" if v["fark"] >= UYUMSUZLUK else "")
                                     for y, v in x["yetkinlik"].items() if v["ort"] is not None))
    satirlar = [f"{harita[n['aday']]} | {harita[n['mulakatci']]} | {n['yetkinlik']} | puan {n['puan'] if n['puan'] is not None else '-'} | "
                + ("[DEĞERLENDİRME DIŞI] " if n["bulgular"] else "") + gizle(n["not"], harita) for n in notlar]
    mesaj = "\n".join(["<puan_ozeti>", *po, "</puan_ozeti>", "<notlar>", *satirlar, "</notlar>"])
    yanit = llm.json_iste(sistem, mesaj, SEMA)
    geri = {v: k for k, v in harita.items()}
    return {geri[s["aday"]]: s for s in yanit.get("adaylar", []) if s.get("aday") in geri}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
ONAY_DOLGU = PatternFill("solid", fgColor="FFF4CE")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
UST = Alignment(vertical="top", wrap_text=True)
ONERI = {"ilerlet": "İlerlet", "beklet": "Beklet", "ilerletme": "İlerletme"}


def rapor_yaz(cikti: Path, notlar: list[dict], h: dict, ozet: dict, agirlik: dict, harita: dict) -> None:
    wb = Workbook()
    k = wb.active
    k.title = "Karşılaştırma"
    ys = h["yetkinlikler"]
    k.append(["Sıra", "Aday"] + [f"{y} (ağ. {agirlik.get(y, 1):g})" for y in ys] + ["Toplam (5)", "Toplam — değ. dışı hariç", "Uyumsuz Yetkinlik",
                                                                                      "Eksik Puan", "Öneri (taslak)", "Karar"])
    for c in k[1]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI
    for a in sorted(h["adaylar"], key=lambda a: h["adaylar"][a]["sira"]):
        x = h["adaylar"][a]
        s = ozet.get(a, {})
        k.append([x["sira"], a] + [None if x["yetkinlik"][y]["ort"] is None else round(x["yetkinlik"][y]["ort"], 2) for y in ys]
                 + [None if x["toplam"] is None else round(x["toplam"], 2), None if x["toplam_temiz"] is None else round(x["toplam_temiz"], 2),
                    ", ".join(f"{y} ({f:g})" for y, f, _ in x["uyumsuz"]), len(x["eksik"]), ONERI.get(s.get("oneri"), ""), ""])
        r = k.max_row
        for j, y in enumerate(ys, 3):
            if x["yetkinlik"][y]["fark"] >= UYUMSUZLUK:
                k.cell(r, j).fill = KIRMIZI
        if x["onyargi"]:
            k.cell(r, 4 + len(ys)).fill = KIRMIZI
        k.cell(r, 7 + len(ys)).fill = AI_DOLGU
        k.cell(r, 8 + len(ys)).fill = ONAY_DOLGU
    for j, w in enumerate([6, 22] + [16] * len(ys) + [11, 14, 30, 10, 14, 12], 1):
        k.column_dimensions[get_column_letter(j)].width = w
    k.append([])
    k.append(["", f"Uyumsuz: mülakatçılar arası puan farkı ≥ {UYUMSUZLUK} (kalibrasyon toplantısında konuşulmalı). "
                  "'Değerlendirme dışı hariç' toplam, korunan özelliğe dayalı yorum içeren notların puanlarını dışarıda bırakır."])

    o = wb.create_sheet("Aday Özetleri")
    o.append(["Aday", "Güçlü Yönler", "Gelişim Alanları", "Görüş Ayrılıkları", "Doğrulanacaklar", "Öneri", "Gerekçe"])
    for c in o[1]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI
    for a in sorted(h["adaylar"], key=lambda a: h["adaylar"][a]["sira"]):
        s = ozet.get(a)
        if not s:
            o.append([a, "AI yanıt vermedi"])
            continue
        g = lambda liste: "\n".join(f"• {ac(x, harita)}" for x in liste)  # noqa: E731
        o.append([a, g(s["guclu_yonler"]), g(s["gelisim_alanlari"]), g(s["gorus_ayriliklari"]), g(s["dogrulanacaklar"]), ONERI[s["oneri"]],
                  ac(s["oneri_gerekce"], harita)])
        for c in o[o.max_row]:
            c.alignment = UST
            c.fill = AI_DOLGU
    for j, w in enumerate((18, 50, 50, 44, 40, 10, 50), 1):
        o.column_dimensions[get_column_letter(j)].width = w

    n = wb.create_sheet("Notlar")
    n.append(["Aday", "Mülakatçı", "Yetkinlik", "Puan", "Not", "Uyarı"])
    for c in n[1]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI
    for x in notlar:
        uyari = ("DEĞERLENDİRME DIŞI — korunan özelliğe dayalı yorum: " + ", ".join(f"'{b['ifade']}'" for b in x["bulgular"])) if x["bulgular"] else ""
        n.append([x["aday"], x["mulakatci"], x["yetkinlik"], x["puan"], x["not"], uyari])
        if uyari:
            n.cell(n.max_row, 6).fill = KIRMIZI
        for c in n[n.max_row]:
            c.alignment = UST
    for j, w in enumerate((18, 22, 28, 7, 70, 50), 1):
        n.column_dimensions[get_column_letter(j)].width = w
    n.freeze_panes = "A2"
    n.auto_filter.ref = n.dimensions

    b = wb.create_sheet("Bilgi")
    for s in [["Model", llm.kullanim_ozeti()], ["Oluşturulma", datetime.now().strftime("%d.%m.%Y %H:%M")],
              ["Gizlilik", "Aday ve mülakatçı adları modele takma adla (A1, M1) gönderildi; rapor şirket içi olduğundan gerçek adlar gösterilir"],
              ["Önemli", "Puanlar ve hesaplar koddan, özetler modelden gelir. Öneri taslaktır; işe alım kararı yöneticiye aittir. "
                         "Korunan özelliğe dayalı yorumlar değerlendirmede kullanılmamalı ve mülakatçılarla paylaşılarak düzeltilmelidir."]]:
        b.append(s)
    b.column_dimensions["A"].width = 16
    b.column_dimensions["B"].width = 120
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(girdi: Path, cikti: Path, yetkinlik_yolu: Path | None = None, evet: bool = False) -> dict:
    notlar = notlari_oku(girdi)
    agirlik = agirliklari_oku(yetkinlik_yolu)
    h = hesapla(notlar, agirlik)
    onyargi = sum(len(x["onyargi"]) for x in h["adaylar"].values())
    print(f"[OK] {len(h['adaylar'])} aday · {len(h['mulakatcilar'])} mülakatçı · {len(h['yetkinlikler'])} yetkinlik · "
          f"değerlendirme dışı not: {onyargi}")
    harita = takma(notlar, h)
    llm.onay_al("Mülakat puanları ve notları (aday/mülakatçı adları takma adlı; e-posta/telefon maskeli) özet için gönderilecek.", evet)
    ozet = ozetle(notlar, h, harita)
    rapor_yaz(cikti, notlar, h, ozet, agirlik, harita)
    return {"hesap": h, "ozet": ozet, "harita": harita}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Birden çok mülakatçının puan ve notlarından aday başına karşılaştırmalı özet hazırlar.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "mulakat_notlari.csv",
                   help="Mülakat notları (.xlsx/.csv): Aday, Mülakatçı, Yetkinlik, Puan (1-5), Not")
    p.add_argument("--yetkinlikler", type=Path, help="Yetkinlik ağırlıkları (.xlsx/.csv: Yetkinlik, Ağırlık)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "mulakat_degerlendirme.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    if a.girdi == BURASI / "ornek_veri" / "mulakat_notlari.csv":
        a.yetkinlikler = a.yetkinlikler or BURASI / "ornek_veri" / "yetkinlikler.csv"
    try:
        s = calistir(a.girdi, a.cikti, a.yetkinlikler, a.evet)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    for ad, x in sorted(s["hesap"]["adaylar"].items(), key=lambda i: i[1]["sira"]):
        print(f"[{x['sira']}] {ad}: {sayi_yaz(x['toplam_temiz'])}" + (f" (tüm notlarla {sayi_yaz(x['toplam'])})" if x["onyargi"] else "")
              + (f" · uyumsuz: {', '.join(y for y, _, _ in x['uyumsuz'])}" if x["uyumsuz"] else ""))
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
