"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    adaylar = sorted(set(re.findall(r"^(A\d+) \|", kullanici, re.M)))
    return {"adaylar": [{"aday": a, "guclu_yonler": [f"{a} için M1 somut örnek verdi"], "gelisim_alanlari": ["Tahsilat"],
                         "gorus_ayriliklari": [], "dogrulanacaklar": ["Referans"], "oneri": "ilerlet", "oneri_gerekce": f"{a} hedeflerini aştı"}
                        for a in adaylar] + [{"aday": "A99", "guclu_yonler": [], "gelisim_alanlari": [], "gorus_ayriliklari": [],
                                               "dogrulanacaklar": [], "oneri": "beklet", "oneri_gerekce": ""}]}


class HesapTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.n = agent.notlari_oku(ORNEK / "mulakat_notlari.csv")
        cls.h = agent.hesapla(cls.n, agent.agirliklari_oku(ORNEK / "yetkinlikler.csv"))

    def test_agirlikli_toplam(self):
        d, e = self.h["adaylar"]["Deniz Kurgu"], self.h["adaylar"]["Ege Deneme"]
        # Deniz: (30×4 + 25×4,5 + 20×3,5 + 15×3,5 + 10×2,5) / 100 = 3,80
        self.assertAlmostEqual(d["toplam"], 3.80)
        # Ege: (30×4 + 25×2,5 + 20×4 + 15×2 + 10×4) / 100 = 3,325; "evli" notlu puan hariç SO = 3 → 3,45
        self.assertAlmostEqual(e["toplam"], 3.325)
        self.assertAlmostEqual(e["toplam_temiz"], 3.45)
        self.assertEqual((d["sira"], e["sira"]), (1, 2))

    def test_uyumsuzluk_ve_onyargi(self):
        e = self.h["adaylar"]["Ege Deneme"]
        self.assertEqual([y for y, f, _ in e["uyumsuz"]], ["Müşteri ilişkileri ve ikna"])
        self.assertEqual(e["onyargi"], [("Sonuç odaklılık", "Bölge Satış Müdürü", "evli"),
                                        ("Sonuç odaklılık", "Bölge Satış Müdürü", "çocuklu olduğu için")])
        self.assertEqual(self.h["adaylar"]["Deniz Kurgu"]["onyargi"], [])

    def test_takma_ad(self):
        harita = agent.takma(self.n, self.h)
        self.assertEqual(agent.gizle("Deniz Kurgu ile İK Uzmanı görüştü", harita), "A1 ile M2 görüştü")
        self.assertEqual(agent.ac("A1, M2'ye göre", harita), "Deniz Kurgu, İK Uzmanı'ye göre")


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_uctan_uca(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "m.xlsx"
            s = agent.calistir(ORNEK / "mulakat_notlari.csv", cikti, ORNEK / "yetkinlikler.csv", evet=True)
            m = MESAJLAR[0]
            for ad in ("Deniz Kurgu", "Ege Deneme", "Bölge Satış Müdürü", "İK Uzmanı"):
                self.assertNotIn(ad, m)                                    # adlar modele gitmez
            self.assertIn("[DEĞERLENDİRME DIŞI]", m)
            self.assertEqual(set(s["ozet"]), {"Deniz Kurgu", "Ege Deneme"})  # uydurma A99 yok sayılır
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Karşılaştırma", "Aday Özetleri", "Notlar", "Bilgi"])
            ozet = wb["Aday Özetleri"]
            self.assertIn("Deniz Kurgu için Bölge Satış Müdürü", ozet.cell(2, 2).value)   # takma adlar geri açılır
            notlar = [r for r in wb["Notlar"].iter_rows(min_row=2, values_only=True) if r[5]]
            self.assertEqual(len(notlar), 1)


if __name__ == "__main__":
    unittest.main()
