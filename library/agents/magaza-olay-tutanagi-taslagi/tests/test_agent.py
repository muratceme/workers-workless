"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import json
import os
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri" / "olay_notlari.txt"
CEVAP = json.loads((Path(__file__).parent / "sahte_cevaplar.json").read_text(encoding="utf-8"))
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    """İlk turda hırsızlık şüphesi özetine 'ürünleri çaldı' yazar; depo olayında olmayan bir not satırına dayanır."""
    MESAJLAR.append(kullanici)
    d = "<duzeltme>" in kullanici
    if "Kasa 2" in kullanici:
        return copy.deepcopy(CEVAP["O1"])
    if "Kozmetik" in kullanici:
        return copy.deepcopy(CEVAP["O2" if d else "O2_hatali"])
    return copy.deepcopy(CEVAP["O3" if d else "O3_hatali"])


def olaylar():
    return agent.olaylari_oku(ORNEK, agent.takma_adlar(["Selin Deneme"]))


class KodTesti(unittest.TestCase):
    def test_okuma_ve_maskeleme(self):
        o = olaylar()
        self.assertEqual([(x["no"], x["zaman"], len(x["satirlar"])) for x in o],
                         [("O1", datetime(2026, 10, 15, 22, 10), 3), ("O2", datetime(2026, 10, 15, 17, 45), 4), ("O3", datetime(2026, 10, 16, 8, 20), 2)])
        self.assertIn("[KİŞİ-1] çalışıyordu", o[0]["satirlar"][1])
        self.assertNotIn("Selin", " ".join(o[0]["satirlar"]))
        self.assertEqual(agent.saatler("17:42-17:46 arası, 9:05"), {"17:42", "17:46", "09:05"})

    def test_suclama(self):
        self.assertEqual([agent.suclama(t) for t in ("Çıkış kapısında alarm çaldı", "Müşteri ürünleri çaldı", "hırsızlık şüphesi", "telefonu çalıyordu",
                                                    "hırsız kaçtı", "alarm çaldıktan sonra kasten çıktı", "personelin zimmetine geçirdiği", "yalan beyan")],
                         [None, "çaldı", None, None, "hırsız", "kasten", "zimmetine", "yalan"])

    def test_denetim_ve_hesap(self):
        o = {x["no"]: x for x in olaylar()}
        self.assertEqual([agent.denetle(o[n], CEVAP[n]) for n in ("O1", "O2", "O3")], [[], [], []])
        self.assertEqual(agent.denetle(o["O2"], CEVAP["O2_hatali"]), ["O2: Özet: suçlayıcı / yargılayıcı ifade ('çaldı'); 'tespit edildi', 'şüphesi', 'görüldü' gibi nesnel dil kullanın"])
        self.assertEqual(agent.denetle(o["O3"], CEVAP["O3_hatali"]), ["O3: dayanak satırlar boş ya da geçersiz (1-2)"])
        c = copy.deepcopy(CEVAP["O1"])
        c["kasa"]["nottaki_fark"] = "250"                                     # farkı hesaplamış
        c["zaman_cizelgesi"].append({"saat": "20:15", "olay": "Kasa kapatıldı"})
        c["ozet"] += " Toplam 3 iade yapıldı."
        s = agent.denetle(o["O1"], c)
        self.assertEqual(len(s), 3)
        self.assertTrue(any("nottaki_fark '250'" in x for x in s) and any("20:15" in x for x in s) and any("olmayan sayı: 3" in x for x in s))
        self.assertEqual(agent.hesap(CEVAP["O1"]), [("Bilgi", "Kasa farkı (kodla): 18.200,00 − 18.450,00 = -250,00 TL (eksik)"),
                                                   ("Hata", "Notta fark 350,00 TL yazıyor, hesaplanan 250,00 TL: sayımı tekrarlayın")])
        self.assertEqual(agent.hesap(CEVAP["O2"]), [("Bilgi", "Etkilenen ürünlerin tutarı (kodla): 3.400,00 TL")])
        c = copy.deepcopy(CEVAP["O2"])
        c["urunler"][0]["tutar"] = "3.700"
        self.assertEqual(agent.hesap(c)[0][0], "Hata")


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ornek_ve_duzeltme(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            self.assertEqual(agent.main(["--evet", "--cikti", str(cikti)]), 0)
            self.assertEqual(len(MESAJLAR), 5)                         # 3 olay + O2 ve O3 düzeltmesi
            for m in MESAJLAR:
                self.assertNotIn("Selin", m)
            ws = load_workbook(cikti)["Olaylar"]
            self.assertEqual([ws.cell(r, 1).value for r in range(2, 5)], ["OT-261015-01", "OT-261015-02", "OT-261016-03"])
            self.assertEqual([ws.cell(2, j).value for j in (4, 6, 7)], ["Kasa farkı", None, "-250,00 TL"])
            self.assertEqual(ws.cell(3, 6).value, "3.400,00 TL")
            t = ws.cell(2, 9).value
            self.assertIn("Selin Deneme isimli kasa personeli", t)               # takma ad tutanakta geri açılır
            self.assertIn("Notta fark 350,00 TL yazıyor, hesaplanan 250,00 TL", t)
            self.assertEqual([x for x in t.splitlines() if x[:2] in {f"{i}." for i in range(1, 9)}],
                             ["1. Olayın özeti", "2. Zaman çizelgesi", "3. Hesap ve kontroller (kod tarafından)", "4. Kanıtlar", "5. Yapılacaklar",
                              "6. Tamamlanması gereken bilgiler"])
            self.assertIn("Tanık (ad soyad, görev, imza)", t)
            self.assertNotIn("çaldı.", ws.cell(3, 9).value.split("Zaman çizelgesi")[0])

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(ORNEK, Path(tmp) / "t.xlsx", ["Selin Deneme"], tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 2)
            self.assertEqual(agent.main(["--girdi", str(Path(tmp) / "yok.txt"), "--evet"]), 1)
            y = Path(tmp) / "n.txt"
            y.write_text("Olay: 40.10.2026 | Kasa\nx\n", encoding="utf-8")
            self.assertEqual(agent.main(["--girdi", str(y), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
