"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import math
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ANKET = agent.ORNEK / "nps_anketi.csv"
MESAJLAR = []
KURAL = [("teslimat", "Teslimat ve lojistik"), ("kargo", "Teslimat ve lojistik"), ("destek", "Müşteri hizmetleri"), ("müşteri hizmet", "Müşteri hizmetleri"),
         ("fatura", "Müşteri hizmetleri"), ("fiyat", "Fiyat"), ("ürün", "Ürün kalitesi"), ("iade", "İade ve değişim"), ("uygulama", "Web sitesi / uygulama"),
         ("web", "Web sitesi / uygulama"), ("kullanım", "Kullanım kolaylığı")]
OLUMSUZ = ("gecikti", "hasarlı", "zor", "çözmeden", "arızalandı", "üç kez", "çöküyor", "yanlış", "kötü")


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append((sistem, kullanici))
    if "<yorumlar>" in kullanici:
        kayitlar = []
        for i, puan, metin in re.findall(r'<yorum id="([^"]+)" puan="(\d+)">\n(.*?)\n</yorum>', kullanici, re.S):
            k = metin.replace("İ", "i").lower()
            temalar = list(dict.fromkeys(t for a, t in KURAL if a in k)) or ["Diğer"]
            duygu = "olumsuz" if any(o in k for o in OLUMSUZ) else ("karma" if " ama " in k or "biraz" in k else "olumlu")
            alinti = "uydurma bir alıntı cümlesi" if i == "N006" else " ".join(metin.split()[:4])
            kayitlar.append({"id": i, "temalar": temalar + ["Uydurma tema"], "duygu": duygu, "alinti": alinti})
        return {"yorumlar": kayitlar + [{"id": "N999", "temalar": ["Diğer"], "duygu": "notr", "alinti": ""}]}
    return {"oneriler": [
        {"baslik": "Teslimat gecikmelerinde bilgilendirme", "sorun": "...", "oneri": "...", "tema": "Teslimat ve lojistik", "yorum_idleri": ["N126", "N012", "N999"]},
        {"baslik": "İade sürecinin kısaltılması", "sorun": "...", "oneri": "...", "tema": "Olmayan tema", "yorum_idleri": ["N081", "N109"]},
        {"baslik": "Dayanaksız öneri", "sorun": "...", "oneri": "...", "tema": "Fiyat", "yorum_idleri": ["X1"]}]}


class PuanTesti(unittest.TestCase):
    def test_nps_formulu(self):
        def y(p):
            return agent.Yanit("x", None, p, {}, {}, "")
        o = agent.nps_hesapla([y(10)] * 50 + [y(8)] * 30 + [y(3)] * 20)
        self.assertAlmostEqual(o["nps"], 30.0)                                 # %50 − %20
        self.assertAlmostEqual(o["hata"], 1.96 * math.sqrt((0.5 + 0.2 - 0.3 ** 2) / 100) * 100)   # ≈ 15,3
        self.assertEqual((o["destekci"], o["pasif"], o["kotuleyen"]), (50, 30, 20))
        self.assertIsNone(agent.nps_hesapla([])["nps"])
        self.assertAlmostEqual(agent.pearson([1, 2, 3, 4, 5, 1, 2, 3, 4, 5], [2, 4, 6, 8, 10, 2, 4, 6, 8, 10]), 1.0)

    def test_yorumsuz(self):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(ANKET, Path(tmp) / "n.xlsx", yorum=False)
        p = s["puan"]
        self.assertEqual(p["genel"]["n"], 178)                                 # 2 geçersiz puan atıldı
        self.assertEqual(sum(o["n"] for o in p["kirilimlar"]["segment"].values()), 178)
        self.assertEqual([o["ad"] for o in p["olcutler"] if o["konum"] == "Öncelikli iyileştirme"], ["Teslimat"])
        self.assertEqual(len(p["aylik"]), 3)
        turler = {(u["tur"], u["kim"]) for u in s["uyarilar"]}
        self.assertTrue({("Geçersiz NPS puanı", ""), ("Düşük NPS", "Segment: Bireysel")} <= turler)
        self.assertEqual(MESAJLAR, [])


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_uctan_uca(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "n.xlsx"
            s = agent.calistir(ANKET, cikti, evet=True)
            self.assertEqual(len(MESAJLAR), 3)                                 # 52 yorum = 2 paket + 1 öneri çağrısı
            self.assertIn("## Yorum analizi", MESAJLAR[0][0])
            self.assertNotIn("## İyileştirme önerileri", MESAJLAR[0][0])
            self.assertIn("## İyileştirme önerileri", MESAJLAR[2][0])
            self.assertNotIn("0532", MESAJLAR[0][1] + MESAJLAR[1][1])           # telefon maskeli
            self.assertNotIn("Bireysel", MESAJLAR[0][1])                        # segment gönderilmez
            y = {x.no: x for x in s["yorumlu"]}
            self.assertEqual(y["N006"].analiz["alinti"], "")
            self.assertIn("Alıntı yorumda bulunamadı; çıkarıldı", y["N006"].kontrol)
            self.assertEqual(y["N126"].analiz["temalar"], ["Teslimat ve lojistik"])     # uydurma tema atıldı
            self.assertTrue(any("çelişiyor" in k for k in y["N011"].kontrol))            # 10 puan + olumsuz yorum
            self.assertEqual([o["baslik"] for o in s["oneriler"]], ["Teslimat gecikmelerinde bilgilendirme", "İade sürecinin kısaltılması"])
            self.assertEqual(s["oneriler"][0]["yorum_idleri"], ["N126", "N012"])
            self.assertEqual(s["oneriler"][1]["tema"], "")
            turler = {u["tur"] for u in s["uyarilar"]}
            self.assertTrue({"Öneri çıkarıldı", "Puan–yorum çelişkisi"} <= turler)
            self.assertEqual(s["temalar"]["İade ve değişim"]["kotuleyen"], 4)
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Özet", "Ölçütler", "Öneriler", "Temalar", "Yorumlar", "Uyarılar"])
            self.assertEqual(wb["Öneriler"]["E2"].value, 2)

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(agent.main(["--yorum-yok", "--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(agent.main(["--yorum-yok", "--min-yanit", "0"]), 2)
            self.assertEqual(agent.main(["--anket", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
