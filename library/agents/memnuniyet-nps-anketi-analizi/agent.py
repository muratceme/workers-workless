"""
Memnuniyet (NPS) Anketi Analizi — Workers / Workless AI Agent
Müşteri Hizmetleri ve Çağrı Merkezi › Müşteri Deneyimi Uzmanı

1. Kod puanları analiz eder:
   - NPS = destekçi (9–10) % − kötüleyen (0–6) %; %95 hata payı = 1,96 × √((p_d + p_k − (p_d − p_k)²) / n) × 100.
   - Segment ve ay bazında NPS (yanıtı --min-yanit altındaki gruplar "yetersiz" olarak işaretlenir).
   - Memnuniyet ölçütleri (1–5): ortalama, üst iki kutu; NPS puanıyla Pearson korelasyonu (önem) ↔ ortalama (performans)
     → "öncelikli iyileştirme" (önemi yüksek, performansı düşük) alanları.
2. Model yorumları (e-posta, telefon, TCKN, IBAN maskeli) sabit tema listesine göre etiketler, duyguyu ve yorumdan
   birebir kısa bir alıntıyı çıkarır. Kod alıntının yorumda geçtiğini doğrular; puanla duygu çelişen yorumları
   (ör. 9–10 puan + olumsuz yorum) işaretler.
3. Model, kodun saydığı tema özetinden iyileştirme önerileri yazar; her öneri yorum kimliklerine dayanmalıdır.
   Geçerli kimliğe dayanmayan öneri rapora alınmaz. Sayılar koddan gelir.
--yorum-yok ile yalnız puan analizi yapılır (API gerekmez).

Kullanım:
    python agent.py                                              # örnek: 180 yanıt, 3 segment, 52 yorum
    python agent.py --anket nps.xlsx --min-yanit 30
    python agent.py --anket nps.xlsx --temalar "Ürün kalitesi" Fiyat Teslimat "Müşteri hizmetleri" --yorum-yok
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import statistics
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import llm

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
PAKET = 40
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
AYLAR = ["Oca", "Şub", "Mar", "Nis", "May", "Haz", "Tem", "Ağu", "Eyl", "Eki", "Kas", "Ara"]
TEMALAR = ["Ürün kalitesi", "Fiyat", "Teslimat ve lojistik", "Müşteri hizmetleri", "Kullanım kolaylığı", "İade ve değişim", "Web sitesi / uygulama", "Diğer"]
DUYGULAR = ["olumlu", "olumsuz", "karma", "notr"]
NPS_ADLARI = ("nps", "tavsiye", "tavsiye puani", "tavsiye etme", "nps puani")
YORUM_ADLARI = ("yorum", "gorus", "aciklama", "neden", "gorus ve oneriler")
SABIT = {"yanit no": "no", "id": "no", "no": "no", "tarih": "tarih", "anket tarihi": "tarih", "segment": "segment", "musteri segmenti": "segment",
         "kanal": "kanal", "bolge": "bolge", "urun": "urun"}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def metin(x) -> str:
    if isinstance(x, float) and x.is_integer():
        return str(int(x))
    return str(x if x is not None else "").strip()


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def sayi(x) -> float | None:
    try:
        return float(str(x).strip().replace(",", ".")) if metin(x) else None
    except ValueError:
        return None


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


# ----------------------------------------------------------------------------
# Veri ve puan analizi (kod)
# ----------------------------------------------------------------------------

@dataclass
class Yanit:
    no: str
    tarih: date | None
    nps: int
    kirilim: dict
    olcutler: dict
    yorum: str
    maskeli: str = ""
    analiz: dict = field(default_factory=dict)
    kontrol: list[str] = field(default_factory=list)

    @property
    def grup(self) -> str:
        return "Destekçi" if self.nps >= 9 else "Pasif" if self.nps >= 7 else "Kötüleyen"


def oku(yol: Path) -> tuple[list[Yanit], list[str], list[dict]]:
    satirlar = tablo_oku(yol)
    b = [katla(c) for c in satirlar[0]]
    nps_j = next((j for j, a in enumerate(b) if a in NPS_ADLARI or a.startswith("nps") or "tavsiye" in a), None)
    if nps_j is None:
        raise ValueError(f"{yol.name}: NPS sütunu bulunamadı (başlıkta 'NPS' veya 'Tavsiye' geçmeli; 0–10 puan)")
    yorum_j = next((j for j, a in enumerate(b) if a in YORUM_ADLARI or a.startswith("yorum") or a.startswith("gorus")), None)
    sabit = {j: SABIT[a] for j, a in enumerate(b) if a in SABIT}
    olcut_j = {j: metin(satirlar[0][j]) for j in range(len(b)) if j not in sabit and j not in (nps_j, yorum_j) and b[j]}
    yanitlar, uy, gecersiz = [], [], 0
    for i, r in enumerate(satirlar[1:], 2):
        al = lambda j: r[j] if j is not None and j < len(r) else None  # noqa: E731
        v = sayi(al(nps_j))
        if v is None or not v.is_integer() or not 0 <= v <= 10:
            gecersiz += 1
            continue
        alan = {ad: metin(al(j)) for j, ad in sabit.items()}
        olcut = {}
        for j, ad in olcut_j.items():
            p = sayi(al(j))
            if p is not None and 1 <= p <= 5:
                olcut[ad] = p
        yanitlar.append(Yanit(alan.get("no") or f"Y{i}", tarih(al(next((j for j, a in sabit.items() if a == "tarih"), None))), int(v),
                              {k: alan[k] for k in ("segment", "kanal", "bolge", "urun") if alan.get(k)}, olcut, metin(al(yorum_j))))
    if gecersiz:
        uy.append({"onem": "Orta", "tur": "Geçersiz NPS puanı", "kim": "", "aciklama": f"{gecersiz} satırda puan boş veya 0–10 tam sayı değil; alınmadı"})
    olcut_adlari = [ad for ad in olcut_j.values() if sum(ad in y.olcutler for y in yanitlar) >= max(5, len(yanitlar) // 10)]
    return yanitlar, olcut_adlari, uy


def nps_hesapla(yanitlar: list[Yanit]) -> dict:
    n = len(yanitlar)
    if not n:
        return {"n": 0, "nps": None, "hata": None, "destekci": 0, "pasif": 0, "kotuleyen": 0}
    d, k = sum(y.nps >= 9 for y in yanitlar), sum(y.nps <= 6 for y in yanitlar)
    pd, pk = d / n, k / n
    se = math.sqrt(max(pd + pk - (pd - pk) ** 2, 0) / n)
    return {"n": n, "nps": (pd - pk) * 100, "hata": 1.96 * se * 100, "destekci": d, "pasif": n - d - k, "kotuleyen": k}


def pearson(x: list[float], y: list[float]) -> float | None:
    if len(x) < 10:
        return None
    mx, my = statistics.fmean(x), statistics.fmean(y)
    sx, sy = math.sqrt(sum((a - mx) ** 2 for a in x)), math.sqrt(sum((b - my) ** 2 for b in y))
    return sum((a - mx) * (b - my) for a, b in zip(x, y)) / (sx * sy) if sx and sy else None


def puan_analizi(yanitlar: list[Yanit], olcut_adlari: list[str], min_yanit: int) -> dict:
    uy = []
    genel = nps_hesapla(yanitlar)
    kirilimlar = {}
    for alan in ("segment", "kanal", "bolge", "urun"):
        degerler = sorted({y.kirilim[alan] for y in yanitlar if alan in y.kirilim})
        if len(degerler) >= 2:
            kirilimlar[alan] = {d: nps_hesapla([y for y in yanitlar if y.kirilim.get(alan) == d]) for d in degerler}
            for d, o in kirilimlar[alan].items():
                o["yeterli"] = o["n"] >= min_yanit
                if o["yeterli"] and o["nps"] < genel["nps"] - 15:
                    uy.append({"onem": "Orta", "tur": "Düşük NPS", "kim": f"{alan.capitalize()}: {d}",
                               "aciklama": f"NPS {o['nps']:+.0f} (genel {genel['nps']:+.0f}, n = {o['n']}, ± {o['hata']:.0f})"})
    aylik = {}
    for ym in sorted({(y.tarih.year, y.tarih.month) for y in yanitlar if y.tarih}):
        aylik[ym] = nps_hesapla([y for y in yanitlar if y.tarih and (y.tarih.year, y.tarih.month) == ym])
        aylik[ym]["yeterli"] = aylik[ym]["n"] >= min_yanit
    olcutler = []
    for ad in olcut_adlari:
        cift = [(y.olcutler[ad], y.nps) for y in yanitlar if ad in y.olcutler]
        p = [a for a, _ in cift]
        olcutler.append({"ad": ad, "n": len(p), "ort": statistics.fmean(p), "ust2": sum(v >= 4 for v in p) / len(p), "alt2": sum(v <= 2 for v in p) / len(p),
                         "r": pearson(p, [float(b) for _, b in cift])})
    if olcutler:
        r_lar = sorted(o["r"] for o in olcutler if o["r"] is not None)
        ort_lar = sorted(o["ort"] for o in olcutler)
        r_med = statistics.median(r_lar) if r_lar else None
        ort_med = statistics.median(ort_lar)
        for o in olcutler:
            if o["r"] is None or r_med is None:
                o["konum"] = "—"
            elif o["r"] >= r_med and o["ort"] < ort_med:
                o["konum"] = "Öncelikli iyileştirme"
            elif o["r"] >= r_med:
                o["konum"] = "Güçlü yön — koruyun"
            elif o["ort"] < ort_med:
                o["konum"] = "İkincil iyileştirme"
            else:
                o["konum"] = "Yeterli"
    if genel["n"] < min_yanit:
        uy.append({"onem": "Orta", "tur": "Az yanıt", "kim": "Genel", "aciklama": f"{genel['n']} yanıt; NPS'nin hata payı geniştir (± {genel['hata']:.0f})"})
    return {"genel": genel, "kirilimlar": kirilimlar, "aylik": aylik, "olcutler": olcutler, "uyarilar": uy}


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

def sema_yorum(temalar):
    return {"type": "object", "properties": {"yorumlar": {"type": "array", "items": {
        "type": "object",
        "properties": {"id": {"type": "string"}, "temalar": {"type": "array", "items": {"type": "string", "enum": temalar}},
                       "duygu": {"type": "string", "enum": DUYGULAR}, "alinti": {"type": "string"}},
        "required": ["id", "temalar", "duygu", "alinti"], "additionalProperties": False}}},
            "required": ["yorumlar"], "additionalProperties": False}


SEMA_ONERI = {"type": "object", "properties": {"oneriler": {"type": "array", "items": {
    "type": "object",
    "properties": {"baslik": {"type": "string"}, "sorun": {"type": "string"}, "oneri": {"type": "string"}, "tema": {"type": "string"},
                   "yorum_idleri": {"type": "array", "items": {"type": "string"}}},
    "required": ["baslik", "sorun", "oneri", "tema", "yorum_idleri"], "additionalProperties": False}}},
              "required": ["oneriler"], "additionalProperties": False}


def _normal(s: str) -> str:
    return " ".join(str(s).replace("İ", "i").replace("I", "ı").lower().split()).strip(" .,;:!?\"'“”‘’")


def alinti_dogru(alinti: str, kaynak: str) -> bool:
    a = _normal(alinti)
    return len(a) >= 8 and a in _normal(kaynak)


def istem(bolum: str) -> str:
    t = (BURASI / "prompt.md").read_text(encoding="utf-8")
    giris, _, kalan = t.partition("## Yorum analizi")
    yorum, _, oneri = kalan.partition("## İyileştirme önerileri")
    return giris + ("## Yorum analizi" + yorum if bolum == "yorum" else "## İyileştirme önerileri" + oneri)


def yorumlari_analiz_et(yorumlu: list[Yanit], temalar: list[str]) -> None:
    sistem = istem("yorum")
    for i in range(0, len(yorumlu), PAKET):
        parca = yorumlu[i:i + PAKET]
        mesaj = "\n".join(["<tema_listesi>" + "; ".join(temalar) + "</tema_listesi>", "<yorumlar>",
                           *[f'<yorum id="{y.no}" puan="{y.nps}">\n{y.maskeli}\n</yorum>' for y in parca], "</yorumlar>"])
        yanit = llm.json_iste(sistem, mesaj, sema_yorum(temalar))
        harita = {y.no: y for y in parca}
        for s in yanit.get("yorumlar", []):
            y = harita.get(s.get("id"))
            if y is None or y.analiz:
                continue
            s["temalar"] = [t for t in dict.fromkeys(s.get("temalar", [])) if t in temalar] or ["Diğer"]
            if s.get("duygu") not in DUYGULAR:
                s["duygu"] = "notr"
            if s.get("alinti") and not alinti_dogru(s["alinti"], y.maskeli):
                y.kontrol.append("Alıntı yorumda bulunamadı; çıkarıldı")
                s["alinti"] = ""
            y.analiz = s
    for y in yorumlu:
        if not y.analiz:
            y.kontrol.append("Model bu yorum için sonuç döndürmedi")
        elif (y.nps >= 9 and y.analiz["duygu"] == "olumsuz") or (y.nps <= 6 and y.analiz["duygu"] == "olumlu"):
            y.kontrol.append(f"Puan ({y.nps}) ile yorum duygusu ({y.analiz['duygu']}) çelişiyor")


def tema_tablosu(yorumlu: list[Yanit], temalar: list[str]) -> dict:
    t = {}
    for ad in temalar:
        lst = [y for y in yorumlu if y.analiz and ad in y.analiz["temalar"]]
        if lst:
            t[ad] = {"toplam": len(lst), "olumsuz": sum(y.analiz["duygu"] in ("olumsuz", "karma") for y in lst), "olumlu": sum(y.analiz["duygu"] == "olumlu" for y in lst),
                     "kotuleyen": sum(y.grup == "Kötüleyen" for y in lst), "pasif": sum(y.grup == "Pasif" for y in lst), "destekci": sum(y.grup == "Destekçi" for y in lst),
                     "ort_nps_puani": statistics.fmean(y.nps for y in lst)}
    return t


def onerileri_al(yorumlu: list[Yanit], temalar_t: dict) -> tuple[list[dict], list[str]]:
    analizli = [y for y in yorumlu if y.analiz]
    if not analizli:
        return [], []
    satirlar = ["<tema_ozeti>"]
    for t, x in temalar_t.items():
        satirlar.append(f'<tema ad="{t}" yorum="{x["toplam"]}" olumsuz="{x["olumsuz"]}" kotuleyen="{x["kotuleyen"]}" pasif="{x["pasif"]}">')
        ornekler = sorted((y for y in analizli if t in y.analiz["temalar"] and y.analiz["duygu"] != "olumlu"), key=lambda y: y.nps)[:12]
        for y in ornekler:
            satirlar.append(f'  <ornek id="{y.no}" puan="{y.nps}">{y.analiz["alinti"] or y.maskeli[:200]}</ornek>')
        satirlar.append("</tema>")
    satirlar.append("</tema_ozeti>")
    yanit = llm.json_iste(istem("oneri"), "\n".join(satirlar), SEMA_ONERI, max_tokens=8000)
    gecerli = {y.no for y in analizli}
    sonuc, atilan = [], []
    for o in yanit.get("oneriler", [])[:8]:
        ids = [i for i in dict.fromkeys(o.get("yorum_idleri", [])) if i in gecerli]
        if not ids:
            atilan.append(o.get("baslik", "?"))
            continue
        o["yorum_idleri"] = ids
        o["tema"] = o.get("tema") if o.get("tema") in temalar_t else ""
        sonuc.append(o)
    return sonuc, atilan


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
AI = PatternFill("solid", fgColor="EDE7FF")
UST = Alignment(vertical="top", wrap_text=True)
KIRILIM_ADI = {"segment": "Segment", "kanal": "Kanal", "bolge": "Bölge", "urun": "Ürün"}


def _baslik(ws, basliklar, genislik=()):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def _nps_satiri(ad, o):
    return [ad, o["n"], None if o["nps"] is None else round(o["nps"], 1), None if o["hata"] is None else round(o["hata"], 1),
            o["destekci"] / o["n"] if o["n"] else None, o["pasif"] / o["n"] if o["n"] else None, o["kotuleyen"] / o["n"] if o["n"] else None,
            "" if o.get("yeterli", True) else "Yanıt yetersiz"]


def rapor_yaz(cikti: Path, s: dict) -> None:
    p = s["puan"]
    wb = Workbook()
    oz = wb.active
    oz.title = "Özet"
    _baslik(oz, ["Grup", "Yanıt", "NPS", "± Hata Payı (%95)", "Destekçi (9–10)", "Pasif (7–8)", "Kötüleyen (0–6)", "Not"], (26, 8, 8, 12, 12, 10, 12, 16))
    def ekle(satir):
        oz.append(satir)
        for j in (5, 6, 7):
            oz.cell(oz.max_row, j).number_format = "0%"
        v = satir[2]
        if v is not None:
            oz.cell(oz.max_row, 3).fill = PatternFill("solid", fgColor="E3F4E1" if v >= 30 else "FFF4CE" if v >= 0 else "FDE2E1")
    ekle(_nps_satiri("Genel", p["genel"]))
    for alan, d in p["kirilimlar"].items():
        oz.append([])
        oz.append([KIRILIM_ADI[alan]])
        oz.cell(oz.max_row, 1).font = Font(bold=True)
        for ad, o in d.items():
            ekle(_nps_satiri(ad, o))
    if p["aylik"]:
        oz.append([])
        oz.append(["Aylık"])
        oz.cell(oz.max_row, 1).font = Font(bold=True)
        for (y, m), o in p["aylik"].items():
            ekle(_nps_satiri(f"{AYLAR[m - 1]} {y}", o))
    oz.append([])
    oz.append(["NPS = destekçi % − kötüleyen %. Hata payı örneklem büyüklüğüne bağlıdır; iki grubun aralıkları kesişiyorsa fark kesin değildir."])
    oz.append([f"Model: {s['model']}"])

    ol = wb.create_sheet("Ölçütler")
    _baslik(ol, ["Ölçüt", "Yanıt", "Ortalama (1–5)", "Üst 2 Kutu", "Alt 2 Kutu", "NPS ile Korelasyon", "Konum"], (28, 8, 12, 10, 10, 12, 26))
    for o in sorted(p["olcutler"], key=lambda o: -(o["r"] or 0)):
        ol.append([o["ad"], o["n"], round(o["ort"], 2), o["ust2"], o["alt2"], None if o["r"] is None else round(o["r"], 2), o["konum"]])
        ol.cell(ol.max_row, 4).number_format = ol.cell(ol.max_row, 5).number_format = "0%"
        if o["konum"] == "Öncelikli iyileştirme":
            ol.cell(ol.max_row, 7).fill = PatternFill("solid", fgColor="FDE2E1")
    ol.append([])
    ol.append(["Korelasyon, ölçütün tavsiye puanıyla birlikte ne kadar değiştiğini gösterir (önem). Medyanın üstünde korelasyon + medyanın altında "
               "ortalama = öncelikli iyileştirme. Korelasyon nedensellik değildir."])

    on = wb.create_sheet("Öneriler")
    _baslik(on, ["Başlık", "Tema", "Sorun", "Öneri", "Dayanak Yorum", "Kötüleyen", "Örnek Alıntılar (doğrulanmış)", "Sorumlu / Karar"], (30, 20, 50, 50, 9, 9, 70, 22))
    yh = {y.no: y for y in s["yorumlu"]}
    for o in s["oneriler"]:
        dayanak = [yh[i] for i in o["yorum_idleri"]]
        alintilar = [f"„{y.analiz['alinti']}” ({y.no}, {y.nps} puan)" for y in dayanak if y.analiz.get("alinti")][:4]
        on.append([o["baslik"], o["tema"], o["sorun"], o["oneri"], len(dayanak), sum(y.grup == "Kötüleyen" for y in dayanak), "\n".join(alintilar), ""])
        for j in (1, 3, 4):
            on.cell(on.max_row, j).fill = AI
        for j in (3, 4, 7):
            on.cell(on.max_row, j).alignment = UST
        on.cell(on.max_row, 8).fill = PatternFill("solid", fgColor="FFF4CE")

    te = wb.create_sheet("Temalar")
    _baslik(te, ["Tema", "Yorum", "Olumlu", "Olumsuz / Karma", "Destekçi", "Pasif", "Kötüleyen", "Ort. Tavsiye Puanı"], (26, 8, 8, 10, 9, 8, 10, 10))
    for t, x in sorted(s["temalar"].items(), key=lambda i: -i[1]["olumsuz"]):
        te.append([t, x["toplam"], x["olumlu"], x["olumsuz"], x["destekci"], x["pasif"], x["kotuleyen"], round(x["ort_nps_puani"], 1)])

    yo = wb.create_sheet("Yorumlar")
    _baslik(yo, ["Yanıt No", "Puan", "Grup", "Segment", "Yorum (maskeli)", "Temalar", "Duygu", "Alıntı", "Kontrol"], (10, 6, 10, 14, 80, 30, 9, 40, 40))
    for y in sorted(s["yorumlu"], key=lambda y: y.nps):
        yo.append([y.no, y.nps, y.grup, y.kirilim.get("segment", ""), y.maskeli, ", ".join(y.analiz.get("temalar", [])), y.analiz.get("duygu", ""),
                   y.analiz.get("alinti", ""), "; ".join(y.kontrol)])
        yo.cell(yo.max_row, 5).alignment = UST
        for j in (6, 7, 8):
            if y.analiz:
                yo.cell(yo.max_row, j).fill = AI
        if y.kontrol:
            yo.cell(yo.max_row, 9).fill = PatternFill("solid", fgColor="FFF4CE")
    yo.auto_filter.ref = f"A1:I{yo.max_row}"

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Kim", "Açıklama"], (9, 26, 24, 100))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["kim"], u["aciklama"]])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(anket: Path, cikti: Path, *, min_yanit: int = 30, temalar: list[str] | None = None, yorum: bool = True, evet: bool = False) -> dict:
    temalar = list(dict.fromkeys((temalar or TEMALAR[:-1]) + ["Diğer"]))
    yanitlar, olcut_adlari, uyarilar = oku(anket)
    if not yanitlar:
        raise llm.LLMHatasi(f"{anket.name}: geçerli yanıt bulunamadı.")
    puan = puan_analizi(yanitlar, olcut_adlari, min_yanit)
    uyarilar += puan["uyarilar"]
    yorumlu = [y for y in yanitlar if len(y.yorum) >= 3]
    for y in yorumlu:
        y.maskeli = llm.maskele(y.yorum)
    oneriler, tema_t, model = [], {}, "kullanılmadı (--yorum-yok)"
    print(f"[OK] {len(yanitlar)} yanıt · NPS {puan['genel']['nps']:+.0f} (± {puan['genel']['hata']:.0f}) · {len(yorumlu)} yorum")
    if yorum and yorumlu:
        llm.onay_al(f"{len(yorumlu)} müşteri yorumu (e-posta, telefon, TCKN ve IBAN maskeli; müşteri adı ve segment gönderilmez) analiz için gönderilecek.", evet)
        yorumlari_analiz_et(yorumlu, temalar)
        tema_t = tema_tablosu(yorumlu, temalar)
        oneriler, atilan = onerileri_al(yorumlu, tema_t)
        for b in atilan:
            uyarilar.append({"onem": "Bilgi", "tur": "Öneri çıkarıldı", "kim": "", "aciklama": f"'{b}' geçerli bir yorum kimliğine dayanmadığı için rapora alınmadı"})
        for y in yorumlu:
            for k in y.kontrol:
                if "çelişiyor" in k:
                    uyarilar.append({"onem": "Bilgi", "tur": "Puan–yorum çelişkisi", "kim": y.no, "aciklama": k + "; puan ölçeği yanlış anlaşılmış olabilir"})
        model = llm.kullanim_ozeti()
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    uyarilar.sort(key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))
    s = {"yanitlar": yanitlar, "yorumlu": yorumlu, "puan": puan, "temalar": tema_t, "oneriler": oneriler, "uyarilar": uyarilar, "model": model}
    rapor_yaz(cikti, s)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="NPS ve memnuniyet anketlerinin puan ve yorumlarını analiz edip iyileştirme önerileri çıkarır.")
    p.add_argument("--anket", type=Path, default=ORNEK / "nps_anketi.csv",
                   help="Yanıt No, Tarih, Segment, Kanal, NPS (0–10), memnuniyet ölçütleri (1–5; sütun adı ölçüt adıdır), Yorum")
    p.add_argument("--min-yanit", type=int, default=30, help="Bir grubun NPS'sinin yorumlanması için en az yanıt (varsayılan 30)")
    p.add_argument("--temalar", nargs="*", help="Yorum temaları (varsayılan: ürün kalitesi, fiyat, teslimat, müşteri hizmetleri...)")
    p.add_argument("--yorum-yok", action="store_true", help="Yorumları modele gönderme; yalnız puan analizi")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "nps_analizi.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    if not a.anket.exists():
        print(f"[X] Dosya bulunamadı: {a.anket}")
        return 1
    if a.min_yanit < 1:
        print("[X] --min-yanit en az 1 olmalı")
        return 2
    try:
        s = calistir(a.anket, a.cikti, min_yanit=a.min_yanit, temalar=a.temalar, yorum=not a.yorum_yok, evet=a.evet)
    except (llm.LLMHatasi, ValueError) as h:
        print(f"[X] {h}")
        return 1
    for o in s["puan"]["olcutler"]:
        if o["konum"] == "Öncelikli iyileştirme":
            print(f"     Öncelikli iyileştirme: {o['ad']} (ortalama {o['ort']:.2f}, korelasyon {o['r']:.2f})".replace(".", ","))
    for o in s["oneriler"]:
        print(f"     • {o['baslik']} ({len(o['yorum_idleri'])} yorum)")
    for u in s["uyarilar"]:
        if u["onem"] != "Bilgi":
            print(f"[!] {u['kim']} · {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    if not a.yorum_yok and s["yorumlu"]:
        print(f"[i] Kullanım: {s['model']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
