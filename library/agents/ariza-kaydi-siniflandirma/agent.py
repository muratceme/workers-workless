"""
Arıza Kaydı Sınıflandırma — Workers / Workless AI Agent
Bakım Onarım › Bakım Teknisyeni

1. Kod serbest metin makine adını makine listesiyle (kod, ad, diğer adlar) eşleştirir.
2. Model her arıza kaydını etiketler: makine, arıza tipi (mekanik, elektrik, hidrolik, pnömatik, otomasyon, kalıp,
   ayar / operatör, diğer), müdahale türü (temizlik, ayar, parça değişimi, onarım, reset, sıkma, dış servis,
   geçici çözüm), değişen parça, olası kök neden ve tekrarını önleme önerisi.
3. Kod denetler: makine listede mi ve kodun kesin eşleştirdiği makineyle aynı mı, metinde parça değişimi yazıyorsa parça
   belirtilmiş mi ve metinde geçiyor mu, alıntılar metinde var mı. Sorun varsa modele geri gönderilir (en fazla --tur).
4. Kod analiz eder: arıza tipi ve makine bazında Pareto (adet, duruş dakikası), makine bazında ortalama duruş (MTTR
   yaklaşığı), --calisma-saat verilirse MTBF; aynı makinede aynı tip arızanın --pencere gün içinde --esik kez (varsayılan
   30 gün, 3 kez) tekrarlaması "tekrarlayan arıza" olarak işaretlenir.

Kullanım:
    python agent.py                                              # örnek kayıtlar + makine listesi
    python agent.py --girdi arizalar.xlsx --makineler makineler.xlsx --calisma-saat 600
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
PAKET = 25
BILINMIYOR = "BİLİNMİYOR"
TIPLER = ["Mekanik", "Elektrik", "Hidrolik", "Pnömatik", "Otomasyon / yazılım (PLC, sensör, sürücü)", "Kalıp / takım", "Ayar / operatör", "Dış etken / diğer"]
MUDAHALE = ["Temizlik", "Ayar", "Parça değişimi", "Onarım", "Reset / yeniden başlatma", "Sıkma / bağlantı", "Dış servis", "Geçici çözüm"]
PARCA_DEGISIM = re.compile(r"(değiştirildi|değişti|yenisi takıldı|yenilendi|yeni(si)? ile değiş)", re.I)
SEMA = {
    "type": "object",
    "properties": {"kayitlar": {"type": "array", "items": {
        "type": "object",
        "properties": {"id": S, "makine_kodu": S, "ariza_tipi": {"type": "string", "enum": TIPLER}, "mudahale": {"type": "array", "items": {"type": "string", "enum": MUDAHALE}},
                       "degisen_parca": S, "belirti": S, "olasi_kok_neden": S, "onleme_onerisi": S, "alinti": S},
        "required": ["id", "makine_kodu", "ariza_tipi", "mudahale", "degisen_parca", "belirti", "olasi_kok_neden", "onleme_onerisi", "alinti"],
        "additionalProperties": False}}},
    "required": ["kayitlar"],
    "additionalProperties": False,
}


def makineleri_oku(yol: Path) -> dict[str, dict]:
    r = kayitlar(yol, {"kod": ("Makine Kodu", "Ekipman Kodu", "Kod"), "ad": ("Makine Adı", "Ekipman Adı", "Ad"), "hat": ("Hat", "Bölüm", "Lokasyon"),
                       "diger": ("Diğer Adlar", "Eşanlamlılar", "Takma Adlar")}, ("kod",))
    return {str(x["kod"]).strip(): {"ad": str(x["ad"]), "hat": str(x["hat"]), "adlar": {katla(a) for a in [x["kod"], x["ad"], *str(x["diger"]).split(",")] if katla(a)}}
            for x in r if x["kod"]}


def arizalari_oku(yol: Path) -> list[dict]:
    r = kayitlar(yol, {"id": ("Kayıt No", "Arıza No", "İş Emri No", "No"), "zaman": ("Tarih", "Başlangıç", "Arıza Tarihi"), "makine": ("Makine", "Ekipman"),
                       "aciklama": ("Açıklama", "Arıza Açıklaması", "Not", "Yapılan İş"), "durus": ("Duruş (dk)", "Duruş", "Duruş Süresi", "Süre (dk)")},
                 ("aciklama",))
    out = []
    for i, x in enumerate(r, 1):
        if not str(x["aciklama"]).strip():
            continue
        try:
            d = float(sayi_norm(x["durus"])) if str(x["durus"]).strip() else 0.0
        except ValueError:
            d = 0.0
        out.append({"id": str(x["id"] or f"AR{i}"), "zaman": zaman_coz(x["zaman"]), "makine": str(x["makine"]), "aciklama": str(x["aciklama"]), "durus": d})
    return out


def kesin_makine(k: dict, makineler: dict) -> str | None:
    m = katla(k["makine"])
    return next((kod for kod, x in makineler.items() if m in x["adlar"]), None) if m else None


def denetle(kayit: list[dict], sonuc: dict, makineler: dict) -> list[str]:
    sorun = []
    for k in kayit:
        s = sonuc.get(k["id"])
        if not s:
            sorun.append(f"{k['id']}: sonuç verilmedi")
            continue
        metin = f"{k['makine']} {k['aciklama']}"
        if s["makine_kodu"] != BILINMIYOR and s["makine_kodu"] not in makineler:
            sorun.append(f"{k['id']}: makine kodu '{s['makine_kodu']}' listede yok (listeden seçin ya da {BILINMIYOR})")
        kesin = kesin_makine(k, makineler)
        if kesin and s["makine_kodu"] != kesin:
            sorun.append(f"{k['id']}: makine alanı '{k['makine']}' listede {kesin} olarak tanımlı, {s['makine_kodu']} verilmiş")
        if PARCA_DEGISIM.search(k["aciklama"]):
            if not s["degisen_parca"]:
                sorun.append(f"{k['id']}: metinde parça değişimi var ama değişen parça boş")
            elif "Parça değişimi" not in s["mudahale"]:
                sorun.append(f"{k['id']}: metinde parça değişimi var ama müdahalede 'Parça değişimi' yok")
        if s["degisen_parca"] and not metinde_var(s["degisen_parca"], metin):
            sorun.append(f"{k['id']}: değişen parça '{s['degisen_parca']}' metinde geçmiyor (metindeki ifadeyi kullanın)")
        if not s["mudahale"]:
            sorun.append(f"{k['id']}: müdahale türü boş")
        if not metinde_var(s["alinti"], metin):
            sorun.append(f"{k['id']}: alıntı metinde yok: '{s['alinti'][:50]}'")
    return sorun


def tekrarlayanlar(kayit: list[dict], sonuc: dict, pencere: int, esik: int) -> list[dict]:
    gr = defaultdict(list)
    for k in kayit:
        s = sonuc.get(k["id"])
        if s and s["makine_kodu"] != BILINMIYOR and k["zaman"]:
            gr[(s["makine_kodu"], s["ariza_tipi"])].append(k)
    out = []
    for (m, t), l in gr.items():
        l.sort(key=lambda x: x["zaman"])
        en_iyi = []
        for i in range(len(l)):
            grup = [x for x in l[i:] if x["zaman"] - l[i]["zaman"] <= timedelta(days=pencere)]
            if len(grup) > len(en_iyi):
                en_iyi = grup
        if len(en_iyi) >= esik:
            out.append({"makine": m, "tip": t, "kayitlar": [x["id"] for x in en_iyi], "bas": en_iyi[0]["zaman"], "bit": en_iyi[-1]["zaman"],
                        "durus": sum(x["durus"] for x in en_iyi)})
    return sorted(out, key=lambda x: -len(x["kayitlar"]))


def makine_ozeti(kayit, sonuc, makineler, calisma_saat: float | None) -> list[list]:
    gr = defaultdict(list)
    for k in kayit:
        if k["id"] in sonuc:
            gr[sonuc[k["id"]]["makine_kodu"]].append(k)
    out = []
    for m, l in sorted(gr.items(), key=lambda x: -sum(k["durus"] for k in x[1])):
        top = sum(k["durus"] for k in l)
        mtbf = round((calisma_saat - top / 60) / len(l), 1) if calisma_saat else None
        out.append([m, makineler.get(m, {}).get("ad", ""), makineler.get(m, {}).get("hat", ""), len(l), top, round(top / len(l), 1), mtbf])
    return out


def uret(kayit: list[dict], makineler: dict, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    liste = "\n".join(f"{kod} | {x['ad']} | {x['hat']}" for kod, x in makineler.items())
    sonuc = {}
    if duzeltme is not None:  # düzeltmede yalnız sorunlu kayıtlar yeniden sorulur
        sorunlu = {s.split(":")[0] for s in duzeltme}
        sonuc.update({a: b for a, b in (onceki or {}).items() if a not in sorunlu})
        kayit = [x for x in kayit if x["id"] in sorunlu]
    for j in range(0, len(kayit), PAKET):
        parca = kayit[j:j + PAKET]
        ids = {k["id"] for k in parca}
        d = [x for x in (duzeltme or []) if x.split(":")[0] in ids]
        if duzeltme is not None and not d:
            sonuc.update({a: b for a, b in (onceki or {}).items() if a in ids})
            continue
        satirlar = []
        for k in parca:
            kesin = kesin_makine(k, makineler)
            satirlar.append(f'<kayit id="{k["id"]}" makine="{k["makine"]}" kesin_eslesme="{kesin or "-"}" durus_dk="{k["durus"]:g}">{llm.maskele(k["aciklama"])}</kayit>')
        mesaj = (f"<makine_listesi>\n{liste}\n</makine_listesi>\n<kayitlar>\n" + "\n".join(satirlar) + "\n</kayitlar>"
                 + duzeltme_blogu(d, [b for a, b in (onceki or {}).items() if a in ids]))
        for s in llm.json_iste(sistem, mesaj, SEMA).get("kayitlar", []):
            if s.get("id") in ids:
                sonuc[s["id"]] = s
    return sonuc


def rapor_yaz(cikti, kayit, sonuc, makineler, tekrar, ozet, kalan, tur) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kayıtlar"
    tekrar_id = {i: f"{t['makine']} · {t['tip']}" for t in tekrar for i in t["kayitlar"]}
    baslik(ws, ["Kayıt No", "Tarih", "Makine (kayıt)", "Makine Kodu", "Eşleşme", "Hat", "Açıklama", "Duruş (dk)", "Arıza Tipi", "Müdahale", "Değişen Parça",
                "Belirti", "Olası Kök Neden", "Önleme Önerisi", "Tekrarlayan", "Onay"],
           (10, 15, 18, 12, 14, 12, 50, 9, 22, 22, 18, 30, 36, 36, 22, 8))
    for k in kayit:
        s = sonuc.get(k["id"])
        if not s:
            ws.append([k["id"], k["zaman"], k["makine"], "", "", "", k["aciklama"], k["durus"], "AI yanıt vermedi"])
            continue
        kesin = kesin_makine(k, makineler)
        es = "Kesin (liste)" if kesin else ("Listede yok" if s["makine_kodu"] == BILINMIYOR else "Model önerisi")
        ws.append([k["id"], k["zaman"], k["makine"], s["makine_kodu"], es, makineler.get(s["makine_kodu"], {}).get("hat", ""), k["aciklama"], k["durus"],
                   s["ariza_tipi"], ", ".join(s["mudahale"]), s["degisen_parca"], s["belirti"], s["olasi_kok_neden"], s["onleme_onerisi"],
                   tekrar_id.get(k["id"], ""), ""])
        ws.cell(ws.max_row, 2).number_format = "DD.MM.YYYY HH:MM"
        boya(ws, range(9, 15), AI_DOLGU)
        if es != "Kesin (liste)":
            boya(ws, [4, 5], SARI)
        if k["id"] in tekrar_id:
            boya(ws, [15], KIRMIZI)
        boya(ws, [16], SARI)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:P{ws.max_row}"
    son_not(ws, "Önemli: Etiketler yapay zekâ önerisidir; bakım sorumlusu onaylamalıdır. Kök neden önerileri doğrulanmamış hipotezdir.", kalan, tur)
    hizala(ws)

    t = wb.create_sheet("Tekrarlayan Arızalar", 0)
    baslik(t, ["Makine", "Arıza Tipi", "Kayıt Sayısı", "İlk", "Son", "Toplam Duruş (dk)", "Kayıtlar", "Öneri"], (10, 30, 11, 15, 15, 14, 30, 60))
    for x in tekrar:
        t.append([x["makine"], x["tip"], len(x["kayitlar"]), x["bas"], x["bit"], x["durus"], ", ".join(x["kayitlar"]),
                  "Kalıcı çözüm için kök neden analizi (5 Neden / balık kılçığı) yapın; önleyici bakım planını gözden geçirin"])
        t.cell(t.max_row, 4).number_format = t.cell(t.max_row, 5).number_format = "DD.MM.YYYY"
    if not tekrar:
        t.append(["—", "Tekrarlayan arıza yok"])
    hizala(t)

    p = wb.create_sheet("Pareto")
    baslik(p, ["Arıza Tipi", "Adet", "Duruş (dk)", "Kümülatif Duruş %"], (40, 8, 12, 18))
    tip = defaultdict(lambda: [0, 0.0])
    for k in kayit:
        if k["id"] in sonuc:
            tip[sonuc[k["id"]]["ariza_tipi"]][0] += 1
            tip[sonuc[k["id"]]["ariza_tipi"]][1] += k["durus"]
    top = sum(v[1] for v in tip.values()) or 1
    kum = 0.0
    for ad, (n, d) in sorted(tip.items(), key=lambda x: -x[1][1]):
        kum += d
        p.append([ad, n, d, round(100 * kum / top, 1)])
    if tip:
        g = BarChart()
        g.title = "Arıza tipine göre duruş (dk)"
        g.add_data(Reference(p, min_col=3, min_row=1, max_row=p.max_row), titles_from_data=True)
        g.set_categories(Reference(p, min_col=1, min_row=2, max_row=p.max_row))
        g.height, g.width = 8, 18
        p.add_chart(g, "F2")
    m = wb.create_sheet("Makine Özeti")
    baslik(m, ["Makine Kodu", "Makine Adı", "Hat", "Arıza Sayısı", "Toplam Duruş (dk)", "Ort. Duruş / MTTR (dk)", "MTBF (saat)"], (12, 26, 14, 11, 15, 18, 12))
    for r in ozet:
        m.append(r)
    m.append([])
    m.append(["MTTR burada arıza başına ortalama duruş süresidir. MTBF = (çalışma saati − toplam duruş saati) / arıza sayısı; yalnız --calisma-saat verildiğinde."])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(girdi: Path, makine_yolu: Path, cikti: Path, pencere: int = 30, esik: int = 3, calisma_saat: float | None = None, tur: int = 2,
             evet: bool = False) -> dict:
    kayit = arizalari_oku(girdi)
    makineler = makineleri_oku(makine_yolu)
    if not kayit:
        raise llm.LLMHatasi(f"{girdi.name}: arıza kaydı yok.")
    print(f"[OK] {len(kayit)} kayıt · {len(makineler)} makine · kesin eşleşen {sum(1 for k in kayit if kesin_makine(k, makineler))}")
    llm.onay_al(f"{len(kayit)} arıza kaydının açıklaması ve makine listesi gönderilecek.", evet)
    sonuc, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(kayit, makineler, d, o), lambda x: denetle(kayit, x, makineler), tur)
    tekrar = tekrarlayanlar(kayit, sonuc, pencere, esik)
    ozet = makine_ozeti(kayit, sonuc, makineler, calisma_saat)
    rapor_yaz(cikti, kayit, sonuc, makineler, tekrar, ozet, kalan, yapilan)
    return {"kayitlar": kayit, "sonuc": sonuc, "tekrar": tekrar, "ozet": ozet, "kalan": kalan, "tur": yapilan}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    p = argparse.ArgumentParser(description="Serbest metin arıza kayıtlarını makine, arıza tipi ve müdahale türüne göre etiketler.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "ariza_kayitlari.csv",
                   help="Arıza kayıtları (.xlsx/.csv): Açıklama; Kayıt No, Tarih, Makine, Duruş (dk) isteğe bağlı")
    p.add_argument("--makineler", type=Path, default=BURASI / "ornek_veri" / "makineler.csv", help="Makine listesi: Makine Kodu, Makine Adı, Hat, Diğer Adlar")
    p.add_argument("--pencere", type=int, default=30, help="Tekrarlayan arıza penceresi, gün (varsayılan 30)")
    p.add_argument("--esik", type=int, default=3, help="Pencerede tekrarlayan sayılacak en az kayıt (varsayılan 3)")
    p.add_argument("--calisma-saat", type=float, help="Dönemdeki makine başına planlı çalışma saati (MTBF için)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "ariza_siniflandirma.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.makineler, a.cikti, a.pencere, a.esik, a.calisma_saat, a.tur, a.evet)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(s['sonuc'])}/{len(s['kayitlar'])} kayıt etiketlendi")
    for t in s["tekrar"]:
        print(f"[!] Tekrarlayan: {t['makine']} · {t['tip']} · {len(t['kayitlar'])} kayıt ({', '.join(t['kayitlar'])})")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
