"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

KAYIT = KLASOR / "ornek_veri" / "ariza_kayitlari.csv"
MAKINE = KLASOR / "ornek_veri" / "makineler.csv"
MESAJLAR = []
H, E, M, O, A, K, P, D = "Hidrolik", "Elektrik", "Mekanik", "Otomasyon / yazılım (PLC, sensör, sürücü)", "Ayar / operatör", "Kalıp / takım", "Pnömatik", "Dış etken / diğer"
T = {  # id: (makine, tip, müdahale, parça, alıntı)
    "AR-101": ("ENJ-01", H, ["Temizlik"], "", "eşanjör filtresi tıkalı çıktı"),
    "AR-102": ("ENJ-02", E, ["Parça değişimi"], "Kalıp ısıtıcı rezistansı", "rezistansı yandı"),
    "AR-103": ("KNV-03", A, ["Ayar"], "", "Bant kaydı"),
    "AR-104": ("KMP-01", P, ["Dış servis", "Parça değişimi"], "ventil", "emniyet ventili kaçırıyor"),
    "AR-105": ("ENJ-01", H, ["Temizlik"], "", "filtre temizlendi"),
    "AR-106": ("ROB-01", P, ["Parça değişimi"], "hortum", "hortum çatlak"),
    "AR-107": ("ENJ-02", K, ["Ayar"], "", "Kalıp ayarı bozuk"),
    "AR-108": ("KNV-03", O, ["Reset / yeniden başlatma"], "", "arıza kodu verdi"),
    "AR-109": ("ENJ-01", H, ["Temizlik"], "", "filtre yine tıkalı"),
    "AR-110": ("ROB-01", O, ["Sıkma / bağlantı"], "", "konnektörü gevşek"),
    "AR-111": (agent.BILINMIYOR, D, ["Onarım"], "", "Forklift lastiği patladı"),
    "AR-112": ("ENJ-02", E, ["Parça değişimi"], "termokupl", "termokupl arızası"),
}


def sonuc(dogru=True) -> dict:
    d = {i: {"id": i, "makine_kodu": m, "ariza_tipi": t, "mudahale": mu, "degisen_parca": p, "belirti": "b", "olasi_kok_neden": "k", "onleme_onerisi": "o",
             "alinti": a} for i, (m, t, mu, p, a) in T.items()}
    if not dogru:
        d["AR-109"]["makine_kodu"] = "ENJ-02"
        d["AR-106"]["degisen_parca"] = ""
        d["AR-111"]["makine_kodu"] = "FRK-01"
    return d


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    d = sonuc("<duzeltme>" in kullanici)
    return {"kayitlar": [d[i] for i in re.findall(r'<kayit id="([^"]+)"', kullanici)]}


class KodTesti(unittest.TestCase):
    def test_okuma_ve_eslesme(self):
        k, m = agent.arizalari_oku(KAYIT), agent.makineleri_oku(MAKINE)
        self.assertEqual((len(k), len(m), k[3]["durus"]), (12, 5, 180.0))
        self.assertEqual([agent.kesin_makine(x, m) for x in k[:4]] + [agent.kesin_makine(k[8], m), agent.kesin_makine(k[10], m)],
                         ["ENJ-01", "ENJ-02", "KNV-03", "KMP-01", "ENJ-01", None])

    def test_denetim_ve_analiz(self):
        k, m = agent.arizalari_oku(KAYIT), agent.makineleri_oku(MAKINE)
        self.assertEqual(agent.denetle(k, sonuc(), m), [])
        s = agent.denetle(k, sonuc(False), m)
        self.assertEqual(len(s), 3)
        self.assertTrue(any("AR-109" in x and "ENJ-01" in x for x in s) and any("AR-106: metinde parça" in x for x in s) and any("FRK-01" in x for x in s))
        d = sonuc()
        d["AR-104"]["mudahale"] = ["Dış servis"]
        d["AR-102"]["degisen_parca"] = "ısıtıcı kartuşu"
        self.assertEqual(len(agent.denetle(k, d, m)), 2)
        t = agent.tekrarlayanlar(k, sonuc(), 30, 3)
        self.assertEqual([(x["makine"], x["kayitlar"], x["durus"]) for x in t], [("ENJ-01", ["AR-101", "AR-105", "AR-109"], 135.0)])
        self.assertEqual(agent.tekrarlayanlar(k, sonuc(), 5, 3), [])
        o = agent.makine_ozeti(k, sonuc(), m, 600)
        self.assertEqual([r[0] for r in o], ["KMP-01", "ENJ-02", "ENJ-01", "ROB-01", "KNV-03", agent.BILINMIYOR])
        self.assertEqual(o[2][3:6], [3, 135.0, 45.0])
        self.assertEqual((o[0][6], o[2][6]), (597.0, 199.2))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "a.xlsx"
            s = agent.calistir(KAYIT, MAKINE, cikti, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            self.assertIn('kesin_eslesme="ENJ-01"', MESAJLAR[0])
            self.assertIn("FRK-01", MESAJLAR[1].split("<duzeltme>")[1])
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Tekrarlayan Arızalar", "Kayıtlar", "Pareto", "Makine Özeti"])
            self.assertEqual(wb["Pareto"].cell(2, 1).value, "Pnömatik")
            self.assertEqual(wb["Kayıtlar"].cell(12, 5).value, "Listede yok")

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(KAYIT, MAKINE, Path(tmp) / "a.xlsx", tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 3)
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--girdi", str(KAYIT.with_name("yok.csv")), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
