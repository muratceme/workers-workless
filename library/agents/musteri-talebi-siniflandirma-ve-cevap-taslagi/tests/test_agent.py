"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    idler = re.findall(r'<talep id="([^"]+)"', kullanici)
    # Model bilerek "hatalı": KVKK ve hukuki risk taleplerini düşük aciliyetli ve insansız işaretliyor
    return {"talepler": [{"id": i, "kategori": "Diğer", "ilgili_ekip": "Müşteri hizmetleri", "aciliyet": "dusuk",
                          "duygu": "notr", "ozet": "özet", "cevap_taslagi": "Merhaba, ...", "kullanilan_bilgi": ["Kargo ve teslimat"],
                          "insan_gerekli": False, "insan_gerekce": ""} for i in idler]
            + [{"id": "T-9999", "kategori": "Diğer", "ilgili_ekip": "Hukuk", "aciliyet": "kritik", "duygu": "notr", "ozet": "",
                "cevap_taslagi": "", "kullanilan_bilgi": [], "insan_gerekli": True, "insan_gerekce": "uydurma"}]}


class OnIslemeTesti(unittest.TestCase):
    def setUp(self):
        self.t = {t.id: t for t in agent.talepleri_oku(ORNEK / "talepler.csv")}
        agent.on_isle(list(self.t.values()))

    def test_siparis_no(self):
        self.assertEqual(self.t["T-1001"].siparis, ["2026100345"])
        self.assertEqual(self.t["T-1003"].siparis, ["2026100299"])
        self.assertEqual(self.t["T-1004"].siparis, ["2026100150"])      # vergi no sipariş sanılmaz

    def test_kural_etiketleri(self):
        self.assertEqual(self.t["T-1005"].etiketler, ["KVKK"])
        self.assertIn("Hukuki risk", self.t["T-1006"].etiketler)
        self.assertEqual(self.t["T-1001"].etiketler, [])                # "hâlâ" tek başına tekrar sayılmaz
        self.assertEqual(self.t["T-1008"].onceki, ["T-1001"])           # aynı gönderici
        self.assertIn("Tekrar yazıyor", self.t["T-1008"].etiketler)

    def test_eml(self):
        with tempfile.TemporaryDirectory() as tmp:
            (Path(tmp) / "a.eml").write_bytes(
                "From: Ali <ali@example.com>\nSubject: =?utf-8?q?=C4=B0ade?=\nDate: Tue, 06 Oct 2026 10:00:00 +0300\n"
                "Content-Type: text/plain; charset=utf-8\n\nÜrünü iade etmek istiyorum.\n\n> eski yazışma\n".encode("utf-8"))
            t = agent.talepleri_oku(Path(tmp))[0]
        self.assertEqual((t.konu, t.mesaj, t.tarih.day), ("İade", "Ürünü iade etmek istiyorum.", 6))


class AgentTesti(unittest.TestCase):
    def setUp(self):
        MESAJLAR.clear()

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_uctan_uca_ve_kod_kurallari(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "m.xlsx"
            s = agent.calistir(ORNEK / "talepler.csv", ORNEK / "bilgi_bankasi.md", cikti, evet=True)
            m = MESAJLAR[0]
            self.assertIn("<bilgi_bankasi>", m)
            self.assertNotIn("0532 111 22 33", m)
            self.assertNotIn("zeynep.k@example.com", m)               # gönderen modele gitmez
            self.assertEqual(len(s["sonuc"]), 8)
            self.assertNotIn("T-9999", s["sonuc"])
            for tid in ("T-1005", "T-1006"):                           # kod kuralı modelin üstünde
                self.assertTrue(s["sonuc"][tid]["insan_gerekli"])
                self.assertEqual(s["sonuc"][tid]["aciliyet"], "yuksek")
            self.assertEqual(s["sonuc"]["T-1008"]["aciliyet"], "normal")   # tekrar eden talep
            self.assertFalse(s["sonuc"]["T-1007"]["insan_gerekli"])

            wb = load_workbook(cikti)
            ws = wb["Talepler"]
            bas = [h.value for h in ws[1]]
            self.assertEqual(bas[-1], "Temsilci Onayı")
            satirlar = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True)}
            self.assertEqual(satirlar["T-1005"][bas.index("Yasal Süre")], "05.11.2026 (KVKK 30 gün)")
            self.assertIn("(kural)", satirlar["T-1006"][bas.index("Yapılacak İşlem")])

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_paketleme(self, _):
        with mock.patch.object(agent, "PAKET", 3), tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(ORNEK / "talepler.csv", ORNEK / "bilgi_bankasi.md", Path(tmp) / "m.xlsx", evet=True)
        self.assertEqual(len(MESAJLAR), 3)
        self.assertEqual(len(s["sonuc"]), 8)


if __name__ == "__main__":
    unittest.main()
