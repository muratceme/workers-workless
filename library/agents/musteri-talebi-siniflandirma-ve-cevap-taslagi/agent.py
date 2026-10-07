"""
Müşteri Talebi Sınıflandırma ve Cevap Taslağı — Workers / Workless AI Agent
Müşteri Hizmetleri ve Çağrı Merkezi › Müşteri Temsilcisi

1. Talepler (Excel/CSV listesi veya .eml/.txt klasörü) okunur. Kod; sipariş numarasını, aynı göndericinin
   tekrarlanan taleplerini ve kural etiketlerini (KVKK başvurusu, hukuki süreç / şikâyet sitesi tehdidi,
   acil ifadeler) bulur. E-posta, telefon, IBAN ve TCKN maskelenir.
2. Model her talebi kategori, aciliyet, duygu ve ilgili ekibe göre sınıflandırır ve YALNIZCA şirketin bilgi
   bankasına dayanarak cevap taslağı yazar; bilgi yoksa veya işlem gerekiyorsa 'insan gerekli' işaretler.
3. Kod kuralları modelin üstündedir: KVKK ve hukuki risk etiketli talepler her zaman insana yönlendirilir ve
   en az 'yüksek' aciliyet alır; KVKK başvurularına 30 günlük yasal cevap süresi yazılır (KVKK md. 13).
4. Sonuç Excel'e yazılır; her taslak için 'Temsilci Onayı' sütunu vardır. Hiçbir cevap otomatik gönderilmez.

Kullanım:
    python agent.py                                          # örnek taleplerle dener
    python agent.py --girdi talepler.xlsx --bilgi bilgi_bankasi.md
    python agent.py --girdi ./gelen_kutusu --bilgi bilgi_bankasi.md --kategoriler kategoriler.json
"""
from __future__ import annotations

import argparse
import csv
import email
import json
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from email import policy
from email.utils import parsedate_to_datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import llm

BURASI = Path(__file__).resolve().parent
PAKET = 20                      # bir istekte en fazla talep
MESAJ_SINIRI = 4000             # tek talepte gönderilen en fazla karakter

VARSAYILAN_KATEGORILER = {
    "kategoriler": ["Sipariş ve kargo", "İade ve cayma", "Değişim", "Ayıplı / hasarlı ürün", "Fatura",
                    "Ödeme ve para iadesi", "Ürün bilgisi", "Şikâyet", "KVKK başvurusu", "Teşekkür / öneri", "Diğer"],
    "ekipler": ["Müşteri hizmetleri", "Lojistik", "İade ve değişim", "Kalite", "Muhasebe", "KVKK birimi", "Hukuk"],
}
ACILIYET = ["dusuk", "normal", "yuksek", "kritik"]
DUYGU = ["olumlu", "notr", "olumsuz", "ofkeli"]

# Kod kuralları: (etiket, desen). Kelime başında, büyük/küçük harf duyarsız aranır.
KURALLAR = [
    ("KVKK", re.compile(r"\b(kvkk|6698|kişisel veri|verilerimin|verilerimi sil|unutulma hakkı)", re.I)),
    ("Hukuki risk", re.compile(r"\b(avukat|dava\b|davalık|mahkeme|savcılı|ihtarname|tüketici hakem|hakem heyeti|"
                                r"şikayetvar|şikâyetvar|cimer|tüketici derneği|suç duyurusu)", re.I)),
    ("Acil", re.compile(r"\b(acil|hemen|derhal|bugün|ivedi)", re.I)),
    ("Tekrar yazıyor", re.compile(r"\b(tekrar yazıyorum|yine yazıyorum|üçüncü kez|ikinci kez|kaçıncı kez)", re.I)),
]
SIPARIS = re.compile(r"(?:sipariş\s*(?:no|numarası|numaralı)?\s*[:#]?\s*|#)(\d{6,14})|\b(\d{8,14})\s*numaralı", re.I)


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


@dataclass
class Talep:
    id: str
    tarih: datetime | None
    kanal: str
    gonderen: str
    konu: str
    mesaj: str
    siparis: list[str] = field(default_factory=list)
    etiketler: list[str] = field(default_factory=list)
    onceki: list[str] = field(default_factory=list)


# ----------------------------------------------------------------------------
# Okuma
# ----------------------------------------------------------------------------

def tarih_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%d.%m.%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(x).strip(), f)
        except ValueError:
            pass
    return None


ALANLAR = {
    "id": ("talep no", "talep id", "id", "ticket", "kayıt no", "no"),
    "tarih": ("tarih", "geliş tarihi", "oluşturma tarihi"),
    "kanal": ("kanal", "kaynak"),
    "gonderen": ("gönderen", "müşteri", "e-posta", "kimden"),
    "konu": ("konu", "başlık"),
    "mesaj": ("mesaj", "içerik", "açıklama", "metin", "talep"),
}


def tablo_oku(yol: Path) -> list[Talep]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    s = [r for r in s if any(c not in (None, "") for c in r)]
    b = [kucuk(x) for x in s[0]]
    k = {alan: next((i for i, x in enumerate(b) if x in adlar), None) for alan, adlar in ALANLAR.items()}
    if k["mesaj"] is None:
        raise SystemExit(f"Talep dosyasında 'Mesaj' sütunu bulunamadı. Başlıklar: {s[0]}")
    al = lambda r, a: r[k[a]] if k[a] is not None and k[a] < len(r) else None  # noqa: E731
    return [Talep(str(al(r, "id") or f"T{i}"), tarih_coz(al(r, "tarih")), str(al(r, "kanal") or ""), str(al(r, "gonderen") or ""),
                  str(al(r, "konu") or ""), str(al(r, "mesaj") or "").strip())
            for i, r in enumerate(s[1:], 1) if str(al(r, "mesaj") or "").strip()]


def eml_oku(yol: Path) -> Talep:
    m = email.message_from_bytes(yol.read_bytes(), policy=policy.default)
    govde = m.get_body(preferencelist=("plain", "html"))
    metin = govde.get_content() if govde else ""
    if govde is not None and govde.get_content_type() == "text/html":
        metin = re.sub(r"<[^>]+>", " ", metin)
    # alıntılanan eski yazışmaları at
    metin = re.split(r"\n(?:>|On .+ wrote:|.+ tarihinde .+ yazdı:)", metin)[0]
    tarih = None
    try:
        tarih = parsedate_to_datetime(m["date"]).replace(tzinfo=None) if m["date"] else None
    except (TypeError, ValueError):
        pass
    return Talep(yol.stem, tarih, "E-posta", str(m["from"] or ""), str(m["subject"] or ""), metin.strip())


def talepleri_oku(girdi: Path) -> list[Talep]:
    if girdi.is_dir():
        talepler = []
        for yol in sorted(girdi.iterdir()):
            if yol.suffix.lower() == ".eml":
                talepler.append(eml_oku(yol))
            elif yol.suffix.lower() == ".txt":
                talepler.append(Talep(yol.stem, None, "Dosya", "", "", yol.read_text(encoding="utf-8", errors="replace").strip()))
        return talepler
    return tablo_oku(girdi)


# ----------------------------------------------------------------------------
# Kod kuralları
# ----------------------------------------------------------------------------

def gonderen_anahtari(g: str) -> str:
    e = re.search(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+", g)
    return (e.group(0) if e else g).strip().lower()


def on_isle(talepler: list[Talep]) -> None:
    gorulen: dict[str, list[str]] = defaultdict(list)
    for t in sorted(talepler, key=lambda t: (t.tarih or datetime.min, t.id)):
        metin = f"{t.konu}\n{t.mesaj}"
        t.siparis = list(dict.fromkeys(a or b for a, b in SIPARIS.findall(metin)))
        t.etiketler = [ad for ad, desen in KURALLAR if desen.search(metin)]
        anahtar = gonderen_anahtari(t.gonderen)
        if anahtar:
            t.onceki = list(gorulen[anahtar])
            gorulen[anahtar].append(t.id)
        if t.onceki and "Tekrar yazıyor" not in t.etiketler:
            t.etiketler.append("Tekrar yazıyor")


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

def sema(kat: dict) -> dict:
    return {
        "type": "object",
        "properties": {
            "talepler": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "id": {"type": "string"},
                        "kategori": {"type": "string", "enum": kat["kategoriler"]},
                        "ilgili_ekip": {"type": "string", "enum": kat["ekipler"]},
                        "aciliyet": {"type": "string", "enum": ACILIYET},
                        "duygu": {"type": "string", "enum": DUYGU},
                        "ozet": {"type": "string"},
                        "cevap_taslagi": {"type": "string"},
                        "kullanilan_bilgi": {"type": "array", "items": {"type": "string"}},
                        "insan_gerekli": {"type": "boolean"},
                        "insan_gerekce": {"type": "string"},
                    },
                    "required": ["id", "kategori", "ilgili_ekip", "aciliyet", "duygu", "ozet", "cevap_taslagi",
                                 "kullanilan_bilgi", "insan_gerekli", "insan_gerekce"],
                    "additionalProperties": False,
                },
            },
        },
        "required": ["talepler"],
        "additionalProperties": False,
    }


def talep_xml(t: Talep) -> str:
    ipucu = []
    if t.siparis:
        ipucu.append("sipariş no: " + ", ".join(t.siparis))
    if t.etiketler:
        ipucu.append("kural etiketleri: " + ", ".join(t.etiketler))
    if t.onceki:
        ipucu.append("aynı göndericinin önceki talepleri: " + ", ".join(t.onceki))
    mesaj = llm.maskele(t.mesaj[:MESAJ_SINIRI])
    return (f'<talep id="{t.id}" kanal="{t.kanal}" tarih="{t.tarih:%d.%m.%Y %H:%M}">\n' if t.tarih else
            f'<talep id="{t.id}" kanal="{t.kanal}">\n') + \
        f"<konu>{llm.maskele(t.konu)}</konu>\n<mesaj>{mesaj}</mesaj>\n<ipuclari>{'; '.join(ipucu) or '-'}</ipuclari>\n</talep>"


def siniflandir(talepler: list[Talep], bilgi: str, kat: dict) -> dict[str, dict]:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    sonuc = {}
    for i in range(0, len(talepler), PAKET):
        parca = talepler[i:i + PAKET]
        mesaj = "\n".join([
            f"<bilgi_bankasi>\n{bilgi}\n</bilgi_bankasi>",
            "<kategoriler>" + "; ".join(kat["kategoriler"]) + "</kategoriler>",
            "<ekipler>" + "; ".join(kat["ekipler"]) + "</ekipler>",
            "<talepler>", *[talep_xml(t) for t in parca], "</talepler>",
        ])
        yanit = llm.json_iste(sistem, mesaj, sema(kat))
        gecerli = {t.id for t in parca}
        for s in yanit.get("talepler", []):
            if s.get("id") in gecerli:
                sonuc[s["id"]] = s
    return sonuc


def kurallari_uygula(t: Talep, s: dict) -> list[str]:
    """Kod kuralları modelin kararının üstündedir. Uygulanan düzeltmeleri döndürür."""
    notlar = []
    if "KVKK" in t.etiketler or "Hukuki risk" in t.etiketler:
        if not s["insan_gerekli"]:
            s["insan_gerekli"] = True
            notlar.append("insan onayı zorunlu (kural)")
        if ACILIYET.index(s["aciliyet"]) < ACILIYET.index("yuksek"):
            s["aciliyet"] = "yuksek"
            notlar.append("aciliyet yükseltildi (kural)")
    if "Tekrar yazıyor" in t.etiketler and s["aciliyet"] == "dusuk":
        s["aciliyet"] = "normal"
        notlar.append("tekrar eden talep (kural)")
    return notlar


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

ACILIYET_AD = {"dusuk": "Düşük", "normal": "Normal", "yuksek": "Yüksek", "kritik": "KRİTİK"}
DUYGU_AD = {"olumlu": "Olumlu", "notr": "Nötr", "olumsuz": "Olumsuz", "ofkeli": "Öfkeli"}
ACILIYET_DOLGU = {"kritik": PatternFill("solid", fgColor="F8C4C1"), "yuksek": PatternFill("solid", fgColor="FDE2E1"),
                  "normal": PatternFill("solid", fgColor="FFF4CE"), "dusuk": PatternFill("solid", fgColor="E3F5E1")}
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
ONAY_DOLGU = PatternFill("solid", fgColor="FFF4CE")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
UST = Alignment(vertical="top", wrap_text=True)


def rapor_yaz(cikti: Path, talepler: list[Talep], sonuc: dict[str, dict], duzeltmeler: dict[str, list[str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Talepler"
    bas = ["Talep No", "Tarih", "Kanal", "Konu", "Mesaj", "Sipariş No", "Kural Etiketleri", "Kategori", "Aciliyet", "Duygu",
           "İlgili Ekip", "Özet", "Cevap Taslağı", "Dayandığı Bilgi", "İnsan Gerekli", "Yapılacak İşlem", "Yasal Süre",
           "Temsilci Onayı"]
    ws.append(bas)
    for h in ws[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for j, w in enumerate((10, 15, 11, 22, 50, 14, 20, 18, 10, 10, 18, 40, 70, 24, 10, 40, 14, 14), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    sira = sorted(talepler, key=lambda t: (-ACILIYET.index(sonuc.get(t.id, {}).get("aciliyet", "normal")), t.tarih or datetime.min))
    for t in sira:
        s = sonuc.get(t.id)
        yasal = ""
        if "KVKK" in t.etiketler or (s and s["kategori"] == "KVKK başvurusu"):
            yasal = f"{(t.tarih or datetime.now()) + timedelta(days=30):%d.%m.%Y} (KVKK 30 gün)"
        if s:
            islem = s["insan_gerekce"]
            if duzeltmeler.get(t.id):
                islem = (islem + " · " if islem else "") + "; ".join(duzeltmeler[t.id])
            ws.append([t.id, t.tarih, t.kanal, t.konu, t.mesaj, ", ".join(t.siparis), ", ".join(t.etiketler), s["kategori"],
                       ACILIYET_AD[s["aciliyet"]], DUYGU_AD[s["duygu"]], s["ilgili_ekip"], s["ozet"], s["cevap_taslagi"],
                       ", ".join(s["kullanilan_bilgi"]), "EVET" if s["insan_gerekli"] else "Hayır", islem, yasal, ""])
            ws.cell(ws.max_row, 9).fill = ACILIYET_DOLGU[s["aciliyet"]]
            for c in (8, 10, 11, 12, 13, 14):
                ws.cell(ws.max_row, c).fill = AI_DOLGU
            if s["insan_gerekli"]:
                ws.cell(ws.max_row, 15).fill = ACILIYET_DOLGU["yuksek"]
        else:
            ws.append([t.id, t.tarih, t.kanal, t.konu, t.mesaj, ", ".join(t.siparis), ", ".join(t.etiketler), "", "", "", "",
                       "AI yanıt vermedi", "", "", "EVET", "Elle değerlendirin", yasal, ""])
        ws.cell(ws.max_row, 2).number_format = "DD.MM.YYYY HH:MM"
        ws.cell(ws.max_row, 18).fill = ONAY_DOLGU
        for h in ws[ws.max_row]:
            h.alignment = UST
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    o = wb.create_sheet("Özet", 0)
    o.append(["Müşteri talepleri özeti"])
    o["A1"].font = Font(bold=True, size=13)
    o.append(["Toplam talep", len(talepler)])
    o.append(["İnsan gerekli", sum(1 for s in sonuc.values() if s["insan_gerekli"]) + (len(talepler) - len(sonuc))])
    o.append(["Model", llm.kullanim_ozeti()])
    o.append(["Oluşturulma", datetime.now().strftime("%d.%m.%Y %H:%M")])
    for baslik, sayac in (("Kategori", Counter(s["kategori"] for s in sonuc.values())),
                          ("Aciliyet", Counter(ACILIYET_AD[s["aciliyet"]] for s in sonuc.values())),
                          ("İlgili ekip", Counter(s["ilgili_ekip"] for s in sonuc.values()))):
        o.append([])
        o.append([baslik, "Adet"])
        for h in o[o.max_row]:
            h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        for ad, n in sayac.most_common():
            o.append([ad, n])
    o.append([])
    o.append(["Önemli", "Cevaplar taslaktır ve otomatik gönderilmez. Kategori, aciliyet ve taslaklar modelden; sipariş no, kural "
                        "etiketleri, tekrar tespiti ve KVKK süresi koddan gelir. Kod kuralları (KVKK ve hukuki risk → insan onayı, "
                        "en az yüksek aciliyet) modelin kararının üstündedir."])
    o.column_dimensions["A"].width = 24
    o.column_dimensions["B"].width = 100
    for satir in o.iter_rows():
        for h in satir:
            h.alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(girdi: Path, bilgi_yolu: Path, cikti: Path, kategori_yolu: Path | None = None, evet: bool = False) -> dict:
    talepler = talepleri_oku(girdi)
    if not talepler:
        raise llm.LLMHatasi(f"{girdi}: okunacak talep yok.")
    kat = dict(VARSAYILAN_KATEGORILER)
    if kategori_yolu:
        kat.update(json.loads(kategori_yolu.read_text(encoding="utf-8")))
    bilgi = bilgi_yolu.read_text(encoding="utf-8")
    on_isle(talepler)
    etiket_sayisi = Counter(e for t in talepler for e in t.etiketler)
    print(f"[OK] {len(talepler)} talep · " + (" · ".join(f"{e}: {n}" for e, n in etiket_sayisi.items()) or "kural etiketi yok"))
    llm.onay_al(f"{len(talepler)} talebin konu ve mesaj metni (e-posta/telefon/IBAN/TCKN maskeli) ve bilgi bankası "
                f"(~{len(bilgi):,} karakter) gönderilecek.".replace(",", "."), evet)
    sonuc = siniflandir(talepler, bilgi, kat)
    duzeltmeler = {t.id: kurallari_uygula(t, sonuc[t.id]) for t in talepler if t.id in sonuc}
    rapor_yaz(cikti, talepler, sonuc, duzeltmeler)
    return {"talepler": talepler, "sonuc": sonuc, "duzeltmeler": duzeltmeler}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Müşteri taleplerini sınıflandırır ve bilgi bankasına dayalı cevap taslağı hazırlar.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "talepler.csv",
                   help="Talep listesi (.xlsx/.csv: Talep No, Tarih, Kanal, Gönderen, Konu, Mesaj) veya .eml/.txt klasörü")
    p.add_argument("--bilgi", type=Path, default=BURASI / "ornek_veri" / "bilgi_bankasi.md", help="Şirket bilgi bankası (.md/.txt)")
    p.add_argument("--kategoriler", type=Path, help="Kategori ve ekip listesi (JSON: {\"kategoriler\": [...], \"ekipler\": [...]})")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "musteri_talepleri.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.bilgi, a.cikti, a.kategoriler, a.evet)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    insan = sum(1 for x in s["sonuc"].values() if x["insan_gerekli"])
    print(f"[OK] {len(s['sonuc'])}/{len(s['talepler'])} talep sınıflandırıldı · insan gerekli: {insan}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
