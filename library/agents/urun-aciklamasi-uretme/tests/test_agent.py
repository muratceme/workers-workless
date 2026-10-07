"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import json
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
KANALLAR = agent.kanallari_oku(ORNEK / "kanallar.json")
MESAJLAR = []
DOLGU = "Yumuşak dokusu ve sade tasarımıyla günlük kullanım için hazırlanmıştır. "


def gecerli_taslak(kod, kanal):
    k = KANALLAR[kanal]
    return {"kanal": kanal, "baslik": f"Örnek Home ürün {kod}"[:k["baslik_max"]],
            "meta_aciklama": "" if k["meta_aciklama_max"] == 0 else "Örnek Home ürünü, sade ve kullanışlı.",
            "aciklama": (DOLGU * 12)[:k["aciklama_min"] + 20], "ozellik_maddeleri": ["Madde"] * k["ozellik_maddesi"],
            "anahtar_kelimeler": ["ev tekstili"]}


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    urunler = []
    for kod, govde in re.findall(r'<urun stok_kodu="([^"]+)">(.*?)</urun>', kullanici, re.S):
        istenen = re.search(rf"{re.escape(kod)}: ([^;\n]+)", kullanici.split("Her ürün için yalnız")[1]).group(1).split(", ")
        kanallar = []
        for kanal in istenen:
            t = gecerli_taslak(kod, kanal)
            if "<duzeltme" not in govde and kod == "NV-101":       # ilk taslakta bilerek hata yap
                t["baslik"] = "Çok uzun bir başlık " * 8 if kanal == "Web sitesi" else "Nevresim Örnek Home"
                t["aciklama"] = t["aciklama"] + " En iyi uyku için 300 TC dokuma."
            kanallar.append(t)
        urunler.append({"stok_kodu": kod, "kanallar": kanallar})
    urunler.append({"stok_kodu": "XX-1", "kanallar": []})
    return {"urunler": urunler}


class DenetimTesti(unittest.TestCase):
    def setUp(self):
        self.urun = agent.urunleri_oku(ORNEK / "urunler.csv")[0]

    def test_okuma(self):
        urunler = agent.urunleri_oku(ORNEK / "urunler.csv")
        self.assertEqual([u["kod"] for u in urunler], ["NV-101", "PK-220", "HV-330", "YS-040"])
        self.assertNotIn("Malzeme", urunler[3]["ozellik"])                 # boş özellik gönderilmez

    def test_uydurma_sayi_ve_iddia(self):
        t = gecerli_taslak("NV-101", "Web sitesi")
        t["aciklama"] += " 200x220 cm nevresim, 300 TC, en iyi kalite, %100 pamuk."
        h = agent.denetle(self.urun, "Web sitesi", KANALLAR["Web sitesi"], t, agent.VARSAYILAN_YASAKLI)
        self.assertIn("ürün verisinde olmayan sayı(lar): 300", h)          # 200, 220, 100 veride var
        self.assertIn("kanıtlanamayan iddia: 'en iyi'", h)

    def test_iddia_veride_varsa_serbest(self):
        u = {"kod": "X", "ozellik": {"Stok Kodu": "X", "Malzeme": "Organik pamuk (GOTS sertifikalı)"}}
        t = gecerli_taslak("X", "Web sitesi")
        t["aciklama"] += " Organik pamuktan üretilmiştir."
        self.assertEqual(agent.denetle(u, "Web sitesi", KANALLAR["Web sitesi"], t, agent.VARSAYILAN_YASAKLI), [])

    def test_sinirlar_ve_marka(self):
        t = gecerli_taslak("NV-101", "Pazaryeri")
        t["baslik"] = "Nevresim Örnek Home"
        t["meta_aciklama"] = "meta"
        h = agent.denetle(self.urun, "Pazaryeri", KANALLAR["Pazaryeri"], t, [])
        self.assertIn("başlık markayla ('Örnek Home') başlamalı", h)
        self.assertIn("bu kanalda meta açıklama boş olmalı", h)


class AgentTesti(unittest.TestCase):
    def setUp(self):
        MESAJLAR.clear()

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "u.xlsx"
            s = agent.calistir(ORNEK / "urunler.csv", ORNEK / "kanallar.json", cikti, evet=True)
            self.assertEqual(len(MESAJLAR), 2)                              # ilk üretim + 1 düzeltme turu
            self.assertIn("<duzeltme", MESAJLAR[1])
            self.assertIn("NV-101", MESAJLAR[1])
            self.assertNotIn("PK-220", MESAJLAR[1])                         # yalnız hatalı ürün geri gider
            self.assertEqual(s["hatalar"], {})
            self.assertEqual(s["gecmis"][("NV-101", "Web sitesi")], 1)
            self.assertNotIn("XX-1", s["taslak"])
            ws = load_workbook(cikti)["İçerikler"]
            self.assertEqual(ws.max_row, 1 + 4 * 2)
            self.assertEqual({r[11] for r in ws.iter_rows(min_row=2, values_only=True)}, {"Uygun"})

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(ORNEK / "urunler.csv", ORNEK / "kanallar.json", Path(tmp) / "u.xlsx", tur=0, evet=True)
        self.assertEqual(len(MESAJLAR), 1)
        self.assertEqual(set(s["hatalar"]), {"NV-101"})                     # kalan hata raporda işaretli
        self.assertTrue(any("300" in h for h in s["hatalar"]["NV-101"]["Web sitesi"]))

    def test_kanal_notu_atlanir(self):
        self.assertNotIn("_not", KANALLAR)
        self.assertIn("_not", json.loads((ORNEK / "kanallar.json").read_text(encoding="utf-8")))


if __name__ == "__main__":
    unittest.main()
