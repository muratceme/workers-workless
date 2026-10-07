"""
Ürün Açıklaması Üretme — Workers / Workless AI Agent
E-ticaret › Kategori ve İçerik Yönetimi › E-ticaret İçerik Uzmanı

Ürün özellik tablosundan her satış kanalı için başlık, meta açıklama, ürün açıklaması, özellik maddeleri ve
anahtar kelimeler üretir. Agent döngüsü:
  1. Model taslağı yazar.
  2. Kod her taslağı denetler: kanal karakter sınırları, başlığın markayla başlaması (kanal kuralıysa),
     ürün verisinde olmayan sayılar (uydurma ölçü/gramaj) ve kanıtlanamayan üstünlük/sağlık iddiaları.
  3. Kural dışı taslaklar hata listesiyle modele geri gönderilir (en fazla --tur kez); kalan hatalar raporda
     işaretlenir. Hiçbir metin otomatik yayımlanmaz; 'Onay' sütunu vardır.

Kullanım:
    python agent.py                                          # örnek ürünlerle dener
    python agent.py --girdi urunler.xlsx --kanallar kanallar.json
    python agent.py --girdi urunler.xlsx --kanallar kanallar.json --yasakli yasakli_ifadeler.txt --tur 3
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import llm

BURASI = Path(__file__).resolve().parent
PAKET = 8                        # bir istekte en fazla ürün
VARSAYILAN_YASAKLI = ["en iyi", "en ucuz", "en kaliteli", "en uygun fiyat", "bir numara", "1 numara", "lider", "rakipsiz",
                      "eşsiz", "mucize", "mucizevi", "tedavi", "şifa", "iyileştirir", "antibakteriyel", "antialerjik",
                      "hipoalerjenik", "organik", "garantili", "ömür boyu", "%100 memnuniyet"]


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


# ----------------------------------------------------------------------------
# Okuma
# ----------------------------------------------------------------------------

def urunleri_oku(yol: Path) -> list[dict]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    s = [r for r in s if any(c not in (None, "") for c in r)]
    basliklar = [str(b or "").strip() for b in s[0]]
    i_kod = next((i for i, b in enumerate(basliklar) if kucuk(b) in ("stok kodu", "ürün kodu", "sku", "barkod", "model kodu")), None)
    if i_kod is None:
        raise SystemExit(f"Ürün tablosunda 'Stok Kodu' sütunu gerekli. Başlıklar: {s[0]}")
    urunler = []
    for r in s[1:]:
        ozellik = {basliklar[i]: str(v).strip() for i, v in enumerate(r) if i < len(basliklar) and basliklar[i] and v not in (None, "")}
        if ozellik.get(basliklar[i_kod]):
            urunler.append({"kod": ozellik[basliklar[i_kod]], "ozellik": ozellik})
    return urunler


def kanallari_oku(yol: Path) -> dict:
    ham = json.loads(yol.read_text(encoding="utf-8"))
    return {ad: k for ad, k in ham.items() if not ad.startswith("_")}


# ----------------------------------------------------------------------------
# Kod denetimi
# ----------------------------------------------------------------------------

_SAYI = re.compile(r"\d+(?:[.,]\d+)?")


def sayilar(metin: str) -> set[str]:
    return {x.replace(",", ".").lstrip("0") or "0" for x in _SAYI.findall(metin)}


def denetle(urun: dict, kanal_adi: str, kanal: dict, t: dict, yasakli: list[str]) -> list[str]:
    hatalar = []
    for alan, sinir in (("baslik", kanal.get("baslik_max")), ("meta_aciklama", kanal.get("meta_aciklama_max"))):
        if sinir is not None and len(t[alan]) > sinir:
            hatalar.append(f"{alan} {len(t[alan])} karakter, sınır {sinir}")
    if kanal.get("meta_aciklama_max") == 0 and t["meta_aciklama"]:
        hatalar.append("bu kanalda meta açıklama boş olmalı")
    n = len(t["aciklama"])
    if kanal.get("aciklama_min") and n < kanal["aciklama_min"]:
        hatalar.append(f"açıklama {n} karakter, en az {kanal['aciklama_min']} olmalı")
    if kanal.get("aciklama_max") and n > kanal["aciklama_max"]:
        hatalar.append(f"açıklama {n} karakter, en fazla {kanal['aciklama_max']} olmalı")
    if kanal.get("ozellik_maddesi") and len(t["ozellik_maddeleri"]) != kanal["ozellik_maddesi"]:
        hatalar.append(f"{len(t['ozellik_maddeleri'])} özellik maddesi var, {kanal['ozellik_maddesi']} olmalı")
    marka = next((v for k, v in urun["ozellik"].items() if kucuk(k) == "marka"), "")
    if marka and "marka" in kucuk(kanal.get("baslik_sirasi", "")).split("+")[0] and not kucuk(t["baslik"]).startswith(kucuk(marka)):
        hatalar.append(f"başlık markayla ('{marka}') başlamalı")
    kaynak = " ".join(urun["ozellik"].values())
    uretilen = " ".join([t["baslik"], t["meta_aciklama"], t["aciklama"], *t["ozellik_maddeleri"]])
    fazla = sorted(sayilar(uretilen) - sayilar(kaynak), key=lambda x: float(x))
    if fazla:
        hatalar.append(f"ürün verisinde olmayan sayı(lar): {', '.join(fazla)}")
    kaynak_k, uretilen_k = kucuk(kaynak), kucuk(uretilen)
    for ifade in yasakli:
        if re.search(r"(?<!\w)" + re.escape(kucuk(ifade)) + r"(?!\w)", uretilen_k) and kucuk(ifade) not in kaynak_k:
            hatalar.append(f"kanıtlanamayan iddia: '{ifade}'")
    return hatalar


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

def sema(kanal_adlari: list[str]) -> dict:
    return {
        "type": "object",
        "properties": {
            "urunler": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "stok_kodu": {"type": "string"},
                        "kanallar": {
                            "type": "array",
                            "items": {
                                "type": "object",
                                "properties": {
                                    "kanal": {"type": "string", "enum": kanal_adlari},
                                    "baslik": {"type": "string"},
                                    "meta_aciklama": {"type": "string"},
                                    "aciklama": {"type": "string"},
                                    "ozellik_maddeleri": {"type": "array", "items": {"type": "string"}},
                                    "anahtar_kelimeler": {"type": "array", "items": {"type": "string"}},
                                },
                                "required": ["kanal", "baslik", "meta_aciklama", "aciklama", "ozellik_maddeleri", "anahtar_kelimeler"],
                                "additionalProperties": False,
                            },
                        },
                    },
                    "required": ["stok_kodu", "kanallar"],
                    "additionalProperties": False,
                },
            },
        },
        "required": ["urunler"],
        "additionalProperties": False,
    }


def urun_xml(u: dict, duzeltme: dict[str, list[str]] | None = None, onceki: dict | None = None) -> str:
    satirlar = [f'<urun stok_kodu="{u["kod"]}">'] + [f"  {k}: {v}" for k, v in u["ozellik"].items()]
    if duzeltme:
        for kanal, hatalar in duzeltme.items():
            satirlar.append(f'  <duzeltme kanal="{kanal}">')
            satirlar.append(f"    önceki taslak: {json.dumps(onceki[kanal], ensure_ascii=False)}")
            satirlar += [f"    hata: {h}" for h in hatalar]
            satirlar.append("  </duzeltme>")
    return "\n".join(satirlar + ["</urun>"])


def uret(urunler: list[dict], kanallar: dict, duzeltmeler: dict | None = None, oncekiler: dict | None = None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    sonuc = {}
    for i in range(0, len(urunler), PAKET):
        parca = urunler[i:i + PAKET]
        istenen = {u["kod"]: (list(duzeltmeler[u["kod"]]) if duzeltmeler else list(kanallar)) for u in parca}
        kanal_bilgi = {ad: kanallar[ad] for ad in dict.fromkeys(k for ks in istenen.values() for k in ks)}
        mesaj = "\n".join([
            "<kanallar>", json.dumps(kanal_bilgi, ensure_ascii=False, indent=1), "</kanallar>",
            "<urunler>",
            *[urun_xml(u, duzeltmeler.get(u["kod"]) if duzeltmeler else None, oncekiler.get(u["kod"]) if oncekiler else None) for u in parca],
            "</urunler>",
            "Her ürün için yalnız şu kanalları yaz: " + "; ".join(f"{k}: {', '.join(v)}" for k, v in istenen.items()),
        ])
        yanit = llm.json_iste(sistem, mesaj, sema(list(kanallar)))
        for s in yanit.get("urunler", []):
            if s.get("stok_kodu") in istenen:
                for t in s["kanallar"]:
                    if t["kanal"] in istenen[s["stok_kodu"]]:
                        sonuc.setdefault(s["stok_kodu"], {})[t["kanal"]] = t
    return sonuc


def calistir(girdi: Path, kanal_yolu: Path, cikti: Path, yasakli_yolu: Path | None = None, tur: int = 2, evet: bool = False) -> dict:
    urunler = urunleri_oku(girdi)
    kanallar = kanallari_oku(kanal_yolu)
    yasakli = VARSAYILAN_YASAKLI if not yasakli_yolu else \
        [x.strip() for x in yasakli_yolu.read_text(encoding="utf-8").splitlines() if x.strip() and not x.startswith("#")]
    print(f"[OK] {len(urunler)} ürün × {len(kanallar)} kanal ({', '.join(kanallar)})")
    llm.onay_al(f"{len(urunler)} ürünün özellik tablosu (kişisel veri içermemeli) metin üretimi için gönderilecek; "
                f"kural dışı taslaklar en fazla {tur} kez düzeltmeye geri gönderilir.", evet)
    taslak = uret(urunler, kanallar)
    gecmis = {}                    # (kod, kanal) → düzeltme turu sayısı
    hatalar = {}
    for tur_no in range(tur + 1):
        hatalar = {}
        for u in urunler:
            for k in kanallar:
                t = taslak.get(u["kod"], {}).get(k)
                h = ["model bu kanal için metin üretmedi"] if t is None else denetle(u, k, kanallar[k], t, yasakli)
                if h:
                    hatalar.setdefault(u["kod"], {})[k] = h
        if not hatalar or tur_no == tur:
            break
        duzeltilecek = [u for u in urunler if u["kod"] in hatalar]
        n = sum(len(v) for v in hatalar.values())
        print(f"[i] Düzeltme turu {tur_no + 1}: {n} taslak kural dışı, modele geri gönderiliyor")
        oncekiler = {kod: {k: taslak.get(kod, {}).get(k, {}) for k in v} for kod, v in hatalar.items()}
        yeni = uret(duzeltilecek, kanallar, hatalar, oncekiler)
        for kod, kanal_taslak in yeni.items():
            for k, t in kanal_taslak.items():
                taslak.setdefault(kod, {})[k] = t
                gecmis[(kod, k)] = tur_no + 1
    rapor_yaz(cikti, urunler, kanallar, taslak, hatalar, gecmis)
    return {"urunler": urunler, "taslak": taslak, "hatalar": hatalar, "gecmis": gecmis}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
ONAY_DOLGU = PatternFill("solid", fgColor="FFF4CE")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)


def rapor_yaz(cikti, urunler, kanallar, taslak, hatalar, gecmis):
    wb = Workbook()
    ws = wb.active
    ws.title = "İçerikler"
    bas = ["Stok Kodu", "Kanal", "Başlık", "Başlık Kr.", "Meta Açıklama", "Meta Kr.", "Açıklama", "Açıklama Kr.",
           "Özellik Maddeleri", "Anahtar Kelimeler", "Düzeltme Turu", "Kontrol", "Onay"]
    ws.append(bas)
    for h in ws[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for j, w in enumerate((11, 13, 50, 8, 45, 8, 80, 9, 50, 40, 9, 40, 10), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for u in urunler:
        for k in kanallar:
            t = taslak.get(u["kod"], {}).get(k)
            h = hatalar.get(u["kod"], {}).get(k, [])
            if t:
                ws.append([u["kod"], k, t["baslik"], len(t["baslik"]), t["meta_aciklama"], len(t["meta_aciklama"]), t["aciklama"],
                           len(t["aciklama"]), "\n".join(f"• {m}" for m in t["ozellik_maddeleri"]), ", ".join(t["anahtar_kelimeler"]),
                           gecmis.get((u["kod"], k), 0), "; ".join(h) or "Uygun", ""])
                for c in (3, 5, 7, 9, 10):
                    ws.cell(ws.max_row, c).fill = AI_DOLGU
            else:
                ws.append([u["kod"], k, "", "", "", "", "", "", "", "", "", "; ".join(h), ""])
            ws.cell(ws.max_row, 12).fill = KIRMIZI if h else YESIL
            ws.cell(ws.max_row, 13).fill = ONAY_DOLGU
            for hc in ws[ws.max_row]:
                hc.alignment = UST
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    v = wb.create_sheet("Ürün Verisi")
    anahtarlar = list(dict.fromkeys(k for u in urunler for k in u["ozellik"]))
    v.append(anahtarlar)
    for h in v[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for u in urunler:
        v.append([u["ozellik"].get(k, "") for k in anahtarlar])

    b = wb.create_sheet("Bilgi")
    uygun = sum(1 for u in urunler for k in kanallar if not hatalar.get(u["kod"], {}).get(k))
    for s in [["Sonuç", f"{uygun}/{len(urunler) * len(kanallar)} içerik tüm kontrollerden geçti"],
              ["Kanal kuralları", json.dumps(kanallar, ensure_ascii=False)],
              ["Kontroller", "Karakter sınırları, özellik maddesi sayısı, başlığın markayla başlaması (kanal kuralıysa), ürün verisinde "
                             "olmayan sayılar, kanıtlanamayan üstünlük/sağlık iddiaları (ürün verisinde geçmiyorsa)"],
              ["Model", llm.kullanim_ozeti()],
              ["Oluşturulma", datetime.now().strftime("%d.%m.%Y %H:%M")],
              ["Önemli", "Metinler yapay zekâ taslağıdır; yayımlamadan önce ürün bilgileriyle karşılaştırıp 'Onay' sütununu doldurun."]]:
        b.append(s)
    b.column_dimensions["A"].width = 18
    b.column_dimensions["B"].width = 120
    for satir in b.iter_rows():
        for h in satir:
            h.alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Ürün özelliklerinden kanal kurallarına uygun başlık ve açıklama üretir.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "urunler.csv", help="Ürün özellik tablosu (.xlsx/.csv), Stok Kodu zorunlu")
    p.add_argument("--kanallar", type=Path, default=BURASI / "ornek_veri" / "kanallar.json", help="Kanal kuralları (JSON)")
    p.add_argument("--yasakli", type=Path, help="Yasaklı ifadeler listesi (.txt, satır başına bir ifade)")
    p.add_argument("--tur", type=int, default=2, help="Kural dışı taslaklar için en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "urun_icerikleri.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.kanallar, a.cikti, a.yasakli, a.tur, a.evet)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    kalan = sum(len(v) for v in s["hatalar"].values())
    print(f"[OK] İçerikler üretildi · kural dışı kalan: {kalan}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
