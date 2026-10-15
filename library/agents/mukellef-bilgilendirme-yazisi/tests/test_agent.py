"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

NOT = KLASOR / "ornek_veri" / "konu_notu.txt"
LISTE = KLASOR / "ornek_veri" / "mukellefler.csv"
IMZA = "Örnek Mali Müşavirlik"
MESAJLAR = []
METIN = ("Sayın {UNVAN},\n\n2026 yılı 3. dönem (Ocak–Eylül) geçici vergi beyannamenizi büromuz hazırlayacaktır. Geçici vergi (yıl içinde, o ana kadarki "
         "kazanç üzerinden peşin ödenen vergi) beyannamesinin verilme ve ödeme süresi 17.11.2026 günü sonuna kadardır.\n\nBeyannameyi hazırlayabilmemiz "
         "için Eylül ayı sonu itibarıyla stok sayım bilginizi ve banka hesap dökümlerinizi en geç 06.11.2026 tarihine kadar büromuza iletmenizi rica "
         "ederiz.\n\nDönem içinde ödediğiniz geçici vergiler, yıllık beyannamede hesaplanan vergiden mahsup edilir (yani düşülür).\n\nAyrıca 7566 sayılı "
         "Kanun ile dördüncü geçici vergi dönemi yeniden getirilmiştir; 2026 yılı 4. dönem beyannamesi 17.02.2027 tarihine kadar verilecektir.\n\n"
         "Saygılarımızla,\n{IMZA}")


def grup(kimlik, dogru=True) -> dict:
    if kimlik == "G3":
        return {"kimlik": "G3", "uygulanir": False, "gerekce": "Basit usulde vergilendirilen mükellefler geçici vergi beyannamesi vermez.",
                "satirlar": [6] if dogru else [], "konu": "", "metin": ""}
    g = {"kimlik": kimlik, "uygulanir": True, "gerekce": "Ticari / serbest meslek kazancı nedeniyle geçici vergi veriyor.", "satirlar": [1, 3],
         "konu": "2026 yılı 3. dönem geçici vergi beyannamesi hakkında", "metin": METIN}
    if not dogru and kimlik == "G1":
        g["metin"] = METIN.replace("\n{IMZA}", "").replace("Sayın {UNVAN},", "{UNVAN}'na,")
    if not dogru and kimlik == "G2":
        g["metin"] = METIN.replace("Saygılarımızla", "VUK 359. madde uyarınca geç verilirse 1.500 TL ceza uygulanabilir; süresinde verilirse ceza "
                                                     "kesilmez.\n\nSaygılarımızla")
    return g


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    ids = re.findall(r"^(G\d+) \|", kullanici, re.M)
    return {"gruplar": [grup(k, "<duzeltme>" in kullanici) for k in ids]}


def veri():
    konu, l = agent.notu_oku(NOT)
    return konu, l, agent.gruplar(agent.mukellefleri_oku(LISTE))


class KodTesti(unittest.TestCase):
    def test_okuma(self):
        konu, l, g = veri()
        self.assertEqual((konu, len(l)), ("2026 yılı 3. dönem geçici vergi beyannamesi", 6))
        self.assertTrue(l[0].startswith("2026 yılı 3. dönem"))
        self.assertEqual([(x["kimlik"], x["tur"], x["sayi"]) for x in g],
                         [("G1", "Kurumlar Vergisi", 2), ("G2", "Gelir Vergisi - Serbest Meslek", 1), ("G3", "Gelir Vergisi - Basit Usul", 1),
                          ("G4", "Gelir Vergisi - Bilanço Esası", 1)])
        self.assertEqual(len(agent.cumleler("2026 yılı 3. dönem beyannamesi 17.11.2026 tarihine kadar verilir. Teşekkürler.")), 2)

    def test_denetim(self):
        konu, l, g = veri()
        dogru = {"gruplar": [grup(x["kimlik"]) for x in g]}
        self.assertEqual(agent.denetle(dogru, g, konu, l, 30), [])
        s = agent.denetle({"gruplar": [grup(x["kimlik"], False) for x in g]}, g, konu, l, 30)
        self.assertEqual(len(s), 6, s)
        self.assertTrue(any(x.startswith("G1") and "{IMZA}" in x for x in s) and any(x.startswith("G1") and "ek getirilmiş" in x for x in s)
                        and any(x.startswith("G2") and "'359'" in x for x in s) and any(x.startswith("G2") and "1500" in x for x in s)
                        and any(x.startswith("G2") and "ceza kesilmez" in x for x in s) and any(x.startswith("G3") and "dayanak" in x for x in s))
        t = {"gruplar": [dict(grup(x["kimlik"]), metin=METIN.replace(" (yani düşülür)", "")) for x in g]}
        self.assertEqual(sum("'mahsup'" in x for x in agent.denetle(t, g, konu, l, 30)), 3)
        self.assertTrue(any("kelimelik cümle" in x for x in agent.denetle(dogru, g, konu, l, 12)))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "m.xlsx"
            s = agent.calistir(NOT, LISTE, cikti, IMZA, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            for gizli in ("Örnek Tekstil", "1234567890", "example"):
                self.assertNotIn(gizli, MESAJLAR[0])
            self.assertEqual(re.findall(r"^(G\d+) \|", MESAJLAR[1], re.M), ["G1", "G2", "G3"])
            hazir = [d for _, _, d in s["yazilar"] if d]
            self.assertEqual(len(hazir), 4)
            self.assertTrue(hazir[0].read_text(encoding="utf-8").startswith("Konu: 2026 yılı 3. dönem geçici vergi beyannamesi hakkında\n\nSayın Örnek Tekstil Ltd. Şti.,"))
            self.assertIn("uygulanmıyor", s["yazilar"][3][1])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Şablonlar", "Yazılar"])
            self.assertEqual(len(s["csv"].read_text(encoding="utf-8-sig").split("Sayın ")) - 1, 4)

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(NOT, LISTE, Path(tmp) / "m.xlsx", IMZA, tur=0, evet=True)
            self.assertEqual((len(s["kalan"]), sum(1 for _, _, d in s["yazilar"] if d)), (6, 1))
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--evet", "--not", str(NOT.with_name("yok.txt"))]), 1)


if __name__ == "__main__":
    unittest.main()
