"""
Mükellef Bilgilendirme Yazısı — Gerisi Bende AI Agent
Beyanname › Serbest Muhasebeci Mali Müşavir

1. Kod SMMM'nin konu notunu (konu satırı + numaralı bilgiler) ve mükellef listesini (Ünvan; Mükellefiyet; E-posta) okur;
   mükellefleri mükellefiyet türüne göre gruplar.
2. Model her mükellefiyet grubu için konunun o gruba uygulanıp uygulanmadığını (dayandığı not satırlarıyla) belirler ve
   uygulanan gruplar için {UNVAN} yer tutuculu, anlaşılır dilde bir bilgilendirme yazısı şablonu yazar. Mükellef adı ve
   e-postası modele gönderilmez.
3. Kod denetler:
   - {UNVAN} ve {IMZA} var mı; bilinmeyen yer tutucu ya da yer tutucuya ek var mı
   - Yazıdaki her tarih, tutar ve sayı notta var mı; kanun / madde / tebliğ atfı notta geçiyor mu
   - Mesleki terim (matrah, tevkifat, mahsup, tarh, tahakkuk …) aynı cümlede açıklanmış mı
   - Cümleler çok uzun mu (--cumle-azami kelime); "ceza kesilmez", "garanti" gibi kesin vaat var mı
   - "Uygulanmaz" kararı bir not satırına dayanıyor mu
   Sorun varsa yalnız o grubun şablonu modele geri gönderilir (en fazla --tur kez).
4. Kod şablonları doldurur; yazılar gönderilmez, mükellef başına .txt ve e-posta birleştirme CSV'si yazılır.

Kullanım:
    python agent.py                                                   # örnek geçici vergi notu
    python agent.py --not konu_notu.txt --mukellefler liste.xlsx --imza "SMMM Ad Soyad · 0212 …"
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
SEMA = {
    "type": "object",
    "properties": {"gruplar": {"type": "array", "items": {"type": "object", "properties": {
        "kimlik": S, "uygulanir": {"type": "boolean"}, "gerekce": S, "satirlar": {"type": "array", "items": {"type": "integer"}}, "konu": S, "metin": S},
        "required": ["kimlik", "uygulanir", "gerekce", "satirlar", "konu", "metin"], "additionalProperties": False}}},
    "required": ["gruplar"],
    "additionalProperties": False,
}
YER_TUTUCULAR = {"{UNVAN}", "{IMZA}"}
TERIMLER = ("matrah", "tevkifat", "mahsup", "tarh", "tahakkuk", "istisna", "muafiyet", "indirimli oran", "gecici vergi", "kurum kazanci", "ihtirazi kayit",
            "pismanlik", "uzlasma", "tecil", "re'sen", "kdv iadesi", "stopaj")
ACIKLAMA = re.compile(r"\(|\byani\b|başka bir deyişle|diğer bir deyişle|anlamına gel|olarak adlandırıl|demektir|ifade eder", re.I)
MADDE = re.compile(r"\b(\d{1,5})\s*sayılı|\b(?:md|m)\.\s*(\d{1,4})|\b(\d{1,4})\s*(?:\.|'?\s*[iıuü]nc[iıuü])?\s*madde|\b(\d{1,4})\s*(?:seri\s*)?no\.?\s*['lı]*\s*tebliğ|"
                   r"\b(?:VUK|GVK|KVK|KDVK)\s*(\d{1,4})", re.I)
VAAT = re.compile(r"\b(ceza kesilmez\w*|ceza (?:almazsınız|çıkmaz)|garanti\w*|kesinlikle (?:iade|indirim)\w*|hiçbir risk\w*|sorun çıkmaz)\b", re.I)


def notu_oku(yol: Path) -> tuple[str, list[str]]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: konu notu bulunamadı.")
    konu, satir = "", []
    for s in belge.metin_oku(yol).splitlines():
        if not s.strip():
            continue
        m = re.match(r"^\s*Konu\s*:\s*(.+)$", s, re.I)
        if m and not konu:
            konu = m.group(1).strip()
        else:
            satir.append(re.sub(r"^\s*(?:\d+[.)]|[-•*])\s*", "", s.strip()))
    if not satir:
        raise llm.LLMHatasi(f"{yol.name}: notta bilgi satırı yok.")
    return konu, satir


def mukellefleri_oku(yol: Path) -> list[dict]:
    r = kayitlar(yol, {"unvan": ("Ünvan", "Unvan", "Ad Soyad / Unvan", "Mükellef"), "tur": ("Mükellefiyet", "Mükellefiyet Türü", "Vergi Türü"),
                       "eposta": ("E-posta", "Eposta", "Mail")}, ("unvan", "tur"))
    return [{"unvan": str(x["unvan"]).strip(), "tur": str(x["tur"]).strip(), "eposta": str(x["eposta"]).strip()} for x in r if str(x["unvan"]).strip()]


def gruplar(mukellefler: list[dict]) -> list[dict]:
    out = {}
    for m in mukellefler:
        anahtar = katla(m["tur"])
        if anahtar not in out:
            out[anahtar] = {"kimlik": f"G{len(out) + 1}", "tur": m["tur"], "sayi": 0}
        out[anahtar]["sayi"] += 1
    return list(out.values())


def cumleler(metin: str) -> list[str]:
    sade = re.sub(r"\d+[.,]\d+(?:[.,]\d+)*", "0", metin)  # tarih / tutar içindeki noktalar cümle bölmesin
    sade = re.sub(r"\b(\d+)\.(?=\s)", r"\1", sade)  # sıra sayısı: "3. dönem"
    return [c.strip() for c in re.split(r"(?<=[.!?])\s+|\n+", sade) if c.strip()]


def denetle(t: dict, grp: list[dict], konu: str, l: list[str], cumle_azami: int) -> list[str]:
    sorun, sab = [], {g["kimlik"]: g for g in t["gruplar"]}
    kaynak = konu + "\n" + "\n".join(l)
    izinli_madde = {x for m in MADDE.finditer(kaynak) for x in m.groups() if x}
    for g in grp:
        ad = f"{g['kimlik']} ({g['tur']})"
        s = sab.get(g["kimlik"])
        if not s:
            sorun.append(f"{ad}: değerlendirilmedi")
            continue
        if not s["satirlar"] or [k for k in s["satirlar"] if not 1 <= k <= len(l)]:
            sorun.append(f"{ad}: karar için dayanak satır yok ya da geçersiz")
        if not s["uygulanir"]:
            continue
        tum = s["konu"] + "\n" + s["metin"]
        for y in sorted(YER_TUTUCULAR - set(re.findall(r"\{[A-Z_]+\}", tum))):
            sorun.append(f"{ad}: {y} yer tutucusu yok")
        bilinmeyen = sorted(set(re.findall(r"\{[^}]*\}", tum)) - YER_TUTUCULAR)
        if bilinmeyen:
            sorun.append(f"{ad}: bilinmeyen yer tutucu: {', '.join(bilinmeyen)}")
        ekli = re.findall(r"\{[A-Z_]+\}['’]?[a-zçğıöşü]+", tum)
        if ekli:
            sorun.append(f"{ad}: yer tutucuya ek getirilmiş ({', '.join(ekli)})")
        sade = re.sub(r"\{[^}]*\}", " ", tum)
        maddeler = {x for m in MADDE.finditer(sade) for x in m.groups() if x}
        for m in sorted(maddeler - izinli_madde):
            sorun.append(f"{ad}: '{m}' numaralı kanun / madde / tebliğ atfı notta yok")
        fazla = sorted(sayilar(sade) - sayilar(kaynak) - maddeler)
        if fazla:
            sorun.append(f"{ad}: notta olmayan tarih / tutar / sayı: {', '.join(fazla)}")
        gorulen = set()  # terim yalnız ilk geçtiği cümlede açıklanmalı
        for c in cumleler(s["metin"]):
            kc = katla(c)
            for terim in TERIMLER:
                if terim not in gorulen and re.search(rf"\b{terim}", kc):
                    gorulen.add(terim)
                    if not ACIKLAMA.search(c):
                        sorun.append(f"{ad}: '{terim}' terimi ilk geçtiği cümlede açıklanmamış: '{c[:60]}…' (parantez ya da 'yani' ile açıklayın)")
            if len(c.split()) > cumle_azami:
                sorun.append(f"{ad}: {len(c.split())} kelimelik cümle (sınır {cumle_azami}): '{c[:50]}…'")
        if VAAT.search(sade):
            sorun.append(f"{ad}: kesin vaat ('{VAAT.search(sade).group(0)}')")
    return sorun


def uret(grp, konu, l, cumle_azami, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    sonuc, hedef = {}, grp
    if duzeltme is not None:  # düzeltmede yalnız sorunlu gruplar yeniden yazılır
        sorunlu = {s.split(" ")[0] for s in duzeltme}
        sonuc.update({g["kimlik"]: g for g in (onceki or {}).get("gruplar", []) if g["kimlik"] not in sorunlu})
        hedef = [g for g in grp if g["kimlik"] in sorunlu]
    ids = {g["kimlik"] for g in hedef}
    notlar = "\n".join(f"{i}. {s}" for i, s in enumerate(l, 1))
    g_satir = "\n".join(f"{g['kimlik']} | {g['tur']} | {g['sayi']} mükellef" for g in hedef)
    mesaj = (f"<konu>{konu}</konu>\n<not>\n{notlar}\n</not>\n<mukellef_gruplari>\n{g_satir}\n</mukellef_gruplari>\n<cumle_azami>{cumle_azami}</cumle_azami>"
             + duzeltme_blogu(duzeltme, [g for g in (onceki or {}).get("gruplar", []) if g["kimlik"] in ids]))
    for g in llm.json_iste(sistem, mesaj, SEMA).get("gruplar", []):
        if g.get("kimlik") in ids:
            sonuc[g["kimlik"]] = g
    return {"gruplar": [sonuc[g["kimlik"]] for g in grp if g["kimlik"] in sonuc]}


def doldur(metin: str, m: dict, imza: str) -> str:
    return metin.replace("{UNVAN}", m["unvan"]).replace("{IMZA}", imza).strip()


def rapor_yaz(cikti: Path, grp, t, mukellefler, yazilar, kalan, tur) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Şablonlar"
    baslik(ws, ["Grup", "Mükellefiyet", "Mükellef", "Uygulanır", "Gerekçe", "Konu", "Şablon", "Kontrol", "Onay"], (6, 26, 9, 10, 40, 36, 80, 30, 8))
    sab = {s["kimlik"]: s for s in t["gruplar"]}
    for g in grp:
        s = sab.get(g["kimlik"], {"uygulanir": None, "gerekce": "", "konu": "", "metin": ""})
        sor = [x.split(": ", 1)[1] for x in kalan if x.startswith(g["kimlik"] + " ")]
        ws.append([g["kimlik"], g["tur"], g["sayi"], {True: "Evet", False: "Hayır"}.get(s["uygulanir"], "?"), s["gerekce"], s["konu"], s["metin"],
                   "; ".join(sor) or "Geçti", ""])
        boya(ws, [5, 6, 7], AI_DOLGU)
        boya(ws, [8], KIRMIZI if sor else YESIL)
        boya(ws, [9], SARI)
    son_not(ws, "Önemli: Yazılar gönderilmedi. Mevzuat bilgisini ve tarihleri resmî kaynaktan teyit edip şablonları onayladıktan sonra gönderin.", kalan, tur)
    hizala(ws)
    y = wb.create_sheet("Yazılar")
    baslik(y, ["Ünvan", "Mükellefiyet", "E-posta", "Durum", "Dosya"], (34, 26, 30, 40, 40))
    for m, durum, dosya in yazilar:
        y.append([m["unvan"], m["tur"], m["eposta"], durum, dosya.name if dosya else ""])
        if not dosya:
            boya(y, [4], SARI)
    hizala(y)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    yol = cikti.with_name(cikti.stem + "_birlestirme.csv")
    with yol.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["E-posta", "Konu", "Metin"])
        for m, durum, dosya in yazilar:
            if dosya and m["eposta"]:
                konu, _, metin = dosya.read_text(encoding="utf-8").partition("\n\n")
                w.writerow([m["eposta"], konu.removeprefix("Konu: "), metin])
    return yol


def calistir(not_yolu: Path, liste_yolu: Path, cikti: Path, imza: str, cumle_azami: int = 30, tur: int = 2, evet: bool = False) -> dict:
    konu, l = notu_oku(not_yolu)
    mukellefler = mukellefleri_oku(liste_yolu)
    grp = gruplar(mukellefler)
    if not grp:
        raise llm.LLMHatasi(f"{liste_yolu.name}: mükellef yok.")
    print(f"[OK] Konu: {konu or '(yok)'} · {len(l)} bilgi satırı · {len(mukellefler)} mükellef · {len(grp)} mükellefiyet grubu")
    llm.onay_al(f"Konu notu ve {len(grp)} mükellefiyet türü gönderilecek (mükellef adı, vergi no ve e-posta gönderilmez).", evet)
    t, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(grp, konu, l, cumle_azami, d, o), lambda x: denetle(x, grp, konu, l, cumle_azami), tur)
    sab = {s["kimlik"]: s for s in t["gruplar"]}
    kimlik = {katla(g["tur"]): g["kimlik"] for g in grp}
    klasor = cikti.with_name(cikti.stem + "_yazilar")
    yazilar = []
    for i, m in enumerate(mukellefler, 1):
        gid = kimlik[katla(m["tur"])]
        s = sab.get(gid)
        if any(x.startswith(gid + " ") for x in kalan) or not s:
            yazilar.append((m, f"Şablon {gid} kontrolden geçmedi", None))
        elif not s["uygulanir"]:
            yazilar.append((m, f"Konu bu mükellefiyete uygulanmıyor: {s['gerekce']}", None))
        else:
            klasor.mkdir(parents=True, exist_ok=True)
            dosya = klasor / f"{i:03d}_{re.sub(r'[^A-Za-z0-9]+', '_', katla(m['unvan']))[:40]}.txt"
            dosya.write_text(f"Konu: {doldur(s['konu'], m, imza)}\n\n{doldur(s['metin'], m, imza)}", encoding="utf-8")
            yazilar.append((m, "Hazır (gönderilmedi)", dosya))
    csv_yol = rapor_yaz(cikti, grp, t, mukellefler, yazilar, kalan, yapilan)
    return {"gruplar": grp, "taslak": t, "yazilar": yazilar, "kalan": kalan, "tur": yapilan, "csv": csv_yol}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    o = BURASI / "ornek_veri"
    p = argparse.ArgumentParser(description="Mevzuat değişikliği ya da beyan sonucunu mükellefe anlaşılır dille açıklayan bilgilendirme yazıları hazırlar (göndermez).")
    p.add_argument("--not", dest="not_", type=Path, default=o / "konu_notu.txt", help="Konu notu (.txt/.docx): 'Konu: …' satırı ve numaralı bilgiler")
    p.add_argument("--mukellefler", type=Path, default=o / "mukellefler.csv", help="Mükellef listesi (.xlsx/.csv): Ünvan, Mükellefiyet; E-posta")
    p.add_argument("--imza", default="Örnek Mali Müşavirlik · SMMM · 0212 000 00 00", help="Yazının imza satırı")
    p.add_argument("--cumle-azami", type=int, default=30, help="Bir cümlede en fazla kelime (varsayılan 30)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "mukellef_yazilari.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.not_, a.mukellefler, a.cikti, a.imza, a.cumle_azami, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    hazir = sum(1 for _, _, d in s["yazilar"] if d)
    print(f"[OK] {hazir}/{len(s['yazilar'])} mükellef için yazı hazır (gönderilmedi)")
    for m, durum, d in s["yazilar"]:
        if not d:
            print(f"[i] {m['unvan']}: {durum}")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()} (+ {s['csv'].name})")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
