"""
Tedarikçi Teklifi Bilgi Çıkarma — Workers / Workless AI Agent
Satın Alma › Satın Alma Uzman Yardımcısı

Tedarikçilerden e-posta, PDF veya Word ile gelen fiyat tekliflerini karşılaştırma tablosuna aktarır.
1. Kod teklif belgelerini ve (varsa) satın alma talep listesini okur; telefon, e-posta, IBAN ve kimlik numaralarını maskeler.
2. Model her teklif için tedarikçiyi, teklif no ve tarihini, kalemleri (talep listesindeki hangi malzemeye karşılık
   geldiği, miktar, birim fiyat, tutar), para birimini, KDV dahil / hariç bilgisini ve ödeme, teslim, nakliye, garanti ve
   geçerlilik ifadelerini belgedeki yazımıyla çıkarır.
3. Kod çıkarımı denetler: metinler belgede aynen geçiyor mu, sayılar belgede yazıyor mu, malzeme kodu talep listesinde var
   mı, aynı kod iki kaleme verilmiş mi. Sorun varsa modele geri gönderilir (en fazla --tur kez).
4. Kod ifadelerden sayıları kendisi hesaplar: "peşin" → 0 gün vade, "5 iş günü" ve "24.10.2026" → takvim günü termin,
   "2 yıl" → 24 ay garanti, "15 gün geçerlidir" → geçerlilik tarihi. Bulgular: satır tutarı = miktar × birim fiyat, talep
   miktarından farklı miktar, teklif verilmeyen kalem, süresi dolmuş teklif, KDV dahil fiyat, belirsiz nakliye / termin.
Çıktının "Teklif Listesi" sayfası "Teklif Karşılaştırma" kod bloğunun girdi biçimindedir.

Kullanım:
    python agent.py                                              # örnek talep ve 3 teklif
    python agent.py --girdi ./teklifler --talep talep.xlsx --tarih 16.10.2026
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
ORNEK_TARIH = date(2026, 10, 16)
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
RENK = {"Hata": "FDE2E1", "Dikkat": "FFF4CE", "Bilgi": "E8F0FE", "Hatalı": "FDE2E1", "Kontrol": "FFF4CE", "Uygun": "E3F5E1"}
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}
YUVARLAMA = Decimal("0.05")
AYLAR = ["ocak", "subat", "mart", "nisan", "mayis", "haziran", "temmuz", "agustos", "eylul", "ekim", "kasim", "aralik"]


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def para(x) -> Decimal | None:
    """'19.000,00' · '3,25' · '2.100' · '200' → Decimal. Okunamazsa None."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1 or (s.count(".") == 1 and len(s.split(".")[1]) == 3):
        s = s.replace(".", "")
    try:
        return Decimal(s) if s not in ("", "-", ".") else None
    except InvalidOperation:
        return None


def belgedeki_sayilar(metin: str) -> set[Decimal]:
    return {v for x in re.findall(r"\d+(?:[.,]\d+)*", metin) if (v := para(x)) is not None}


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    m = re.search(r"\b(\d{1,2})[./-](\d{1,2})[./-](\d{4})\b", s)
    if m:
        try:
            return date(int(m[3]), int(m[2]), int(m[1]))
        except ValueError:
            return None
    m = re.search(r"\b(\d{4})-(\d{2})-(\d{2})\b", s)
    if m:
        return date(int(m[1]), int(m[2]), int(m[3]))
    m = re.search(r"\b(\d{1,2}) (" + "|".join(AYLAR) + r") (\d{4})\b", katla(s))
    if m:
        try:
            return date(int(m[3]), AYLAR.index(m[2]) + 1, int(m[1]))
        except ValueError:
            return None
    return None


def tl(x: Decimal) -> str:
    return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def is_gunu_ekle(d: date, n: int) -> date:
    while n:
        d += timedelta(days=1)
        n -= d.weekday() < 5
    return d


# ----------------------------------------------------------------------------
# İfadelerden sayı (kod hesaplar, model değil)
# ----------------------------------------------------------------------------

def vade_gun(metin: str) -> int | None:
    k = katla(metin)
    if not k:
        return None
    if re.search(r"\b(pesin|nakit|siparisle birlikte|on odeme|avans)\b", k) and not re.search(r"\d+ gun", k):
        return 0
    m = re.search(r"\b(\d+) (gun|ay)\b", k)
    if m:
        return int(m[1]) * (30 if m[2] == "ay" else 1)
    return None


def termin_gun(metin: str, bugun: date) -> tuple[int | None, str]:
    """('5 iş günü', bugün) → (7, 'iş günü takvim gününe çevrildi'). Tarih yazılmışsa bugünden itibaren gün sayısı."""
    k = katla(metin)
    if not k:
        return None, ""
    t = tarih_coz(metin)
    if t:
        return (t - bugun).days, f"{t:%d.%m.%Y} teslim tarihi bugüne ({bugun:%d.%m.%Y}) göre güne çevrildi"
    m = re.search(r"\b(\d+) (is gunu|gun|hafta)\b", k)       # "2-3 hafta" katlanınca "2 3 hafta" olur: üst sınır (3) alınır
    if not m:
        return None, ""
    n = int(m[1])
    if m[2] == "is gunu":
        return (is_gunu_ekle(bugun, n) - bugun).days, f"{n} iş günü takvim gününe çevrildi"
    return n * (7 if m[2] == "hafta" else 1), ""


def garanti_ay(metin: str) -> int | None:
    m = re.search(r"\b(\d+) (ay|yil)\b", katla(metin))
    return None if not m else int(m[1]) * (12 if m[2] == "yil" else 1)


def gecerlilik(metin: str, teklif_tarihi: date | None) -> date | None:
    t = tarih_coz(metin)
    if t:
        return t
    m = re.search(r"\b(\d+) (gun|hafta|ay)\b", katla(metin))
    if m and teklif_tarihi:
        return teklif_tarihi + timedelta(days=int(m[1]) * {"gun": 1, "hafta": 7, "ay": 30}[m[2]])
    return None


# ----------------------------------------------------------------------------
# Girdi
# ----------------------------------------------------------------------------

def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


TALEP_SUTUN = {"kod": ("malzeme kodu", "stok kodu", "kod", "kalem no"), "malzeme": ("malzeme", "malzeme adi", "urun", "aciklama", "tanim"),
               "miktar": ("miktar", "talep miktari"), "birim": ("birim",)}


def talep_oku(yol: Path | None) -> list[dict]:
    if yol is None:
        return []
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: talep listesi bulunamadı.")
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: talep listesi boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in es), None) for a, es in TALEP_SUTUN.items()}
    if idx["kod"] is None or idx["malzeme"] is None:
        raise llm.LLMHatasi(f"{yol.name}: 'Malzeme Kodu' ve 'Malzeme' sütunları gerekli. Başlıklar: {s[0]}")
    al = lambda r, a: "" if idx[a] is None or idx[a] >= len(r) or r[idx[a]] is None else str(r[idx[a]]).strip()  # noqa: E731
    out = [{"kod": al(r, "kod"), "malzeme": al(r, "malzeme"), "miktar": para(al(r, "miktar")), "birim": al(r, "birim")} for r in s[1:]]
    return [x for x in out if x["kod"]]


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

IFADE = ["odeme_metin", "termin_metin", "nakliye_metin", "garanti_metin", "gecerlilik_metin"]
ETIKET = {"tedarikci": "Tedarikçi", "teklif_no": "Teklif no", "teklif_tarihi": "Teklif tarihi", "odeme_metin": "Ödeme ifadesi",
          "termin_metin": "Teslim ifadesi", "nakliye_metin": "Nakliye ifadesi", "garanti_metin": "Garanti ifadesi",
          "gecerlilik_metin": "Geçerlilik ifadesi", "nakliye_tutar": "Nakliye tutarı", "malzeme": "Malzeme", "miktar": "Miktar",
          "birim_fiyat": "Birim fiyat", "tutar": "Tutar"}
KALEM = {"type": "object", "properties": {"talep_kodu": S, "malzeme": S, "miktar": S, "birim": S, "birim_fiyat": S, "tutar": S},
         "required": ["talep_kodu", "malzeme", "miktar", "birim", "birim_fiyat", "tutar"], "additionalProperties": False}
SEMA = {"type": "object", "properties": {
    "teklif_var": {"type": "boolean"}, "tedarikci": S, "teklif_no": S, "teklif_tarihi": S, **{a: S for a in IFADE}, "nakliye_tutar": S,
    "para_birimi": {"type": "string", "enum": ["TL", "USD", "EUR", "GBP", "Diğer"]},
    "kdv": {"type": "string", "enum": ["Hariç", "Dahil", "Belirtilmemiş"]},
    "kalemler": {"type": "array", "items": KALEM}, "teklif_verilmeyen": {"type": "array", "items": S},
    "sartlar": {"type": "array", "items": S}, "notlar": {"type": "array", "items": S}},
    "required": ["teklif_var", "tedarikci", "teklif_no", "teklif_tarihi", *IFADE, "nakliye_tutar", "para_birimi", "kdv", "kalemler",
                 "teklif_verilmeyen", "sartlar", "notlar"], "additionalProperties": False}


def denetle(t: dict, metin: str, talep: list[dict]) -> list[str]:
    """Modelin çıkarımındaki sorunlar (teklifin kendi eksikleri değil). Boş liste = sorunsuz."""
    if not t["teklif_var"]:
        return []
    sorun, sayilar = [], belgedeki_sayilar(metin)
    kodlar = {x["kod"] for x in talep}

    def sayi(on, ad, v):
        if not v.strip():
            return
        d = para(v)
        if d is None:
            sorun.append(f"{on}: {ETIKET[ad]} '{v}' sayı olarak okunamadı; belgedeki sayıyı yazın")
        elif d not in sayilar:
            sorun.append(f"{on}: {ETIKET[ad]} {v} belgede yazmıyor (hesaplamayın; belgede yoksa boş bırakın)")

    for a in ("tedarikci", "teklif_no", "teklif_tarihi", *IFADE):
        if t[a].strip() and not metinde_var(t[a], metin):
            sorun.append(f"Teklif: {ETIKET[a]} '{t[a][:60]}' belgede geçmiyor (belgedeki yazımı aynen kullanın)")
    if not t["tedarikci"].strip():
        sorun.append("Teklif: tedarikçi adı boş; belgedeki firma adını yazın")
    sayi("Teklif", "nakliye_tutar", t["nakliye_tutar"])
    if not t["kalemler"]:
        sorun.append("Teklif: kalem yok; fiyat verilen satırları yazın")
    gorulen = {}
    for j, k in enumerate(t["kalemler"], 1):
        on = f"Kalem {j}"
        if not metinde_var(k["malzeme"], metin):
            sorun.append(f"{on}: Malzeme '{k['malzeme'][:50]}' belgede geçmiyor (tedarikçinin yazımını aynen kullanın)")
        for a in ("miktar", "birim_fiyat", "tutar"):
            sayi(on, a, k[a])
        if not k["birim_fiyat"].strip():
            sorun.append(f"{on}: birim fiyat boş")
        kod = k["talep_kodu"].strip()
        if talep and kod and kod not in kodlar:
            sorun.append(f"{on}: talep kodu '{kod}' talep listesinde yok ({', '.join(sorted(kodlar))}); eşleşme yoksa boş bırakın")
        if kod and kod in gorulen:
            sorun.append(f"{on}: talep kodu {kod} kalem {gorulen[kod]} ile aynı; her talep kalemi en fazla bir kaleme eşleşir")
        if kod:
            gorulen.setdefault(kod, j)
    for kod in t["teklif_verilmeyen"]:
        if kod not in kodlar:
            sorun.append(f"Teklif: teklif verilmeyen '{kod}' talep listesinde bir kod değil")
        elif kod in gorulen:
            sorun.append(f"Teklif: {kod} hem fiyatlı kalemlerde hem teklif verilmeyenlerde")
    return sorun


def degerlendir(t: dict, metin: str, talep: list[dict], bugun: date) -> dict:
    """Kod: ifadelerden sayıları hesaplar ve teklifle ilgili bulguları çıkarır."""
    b = []
    tt = tarih_coz(t["teklif_tarihi"])
    v = vade_gun(t["odeme_metin"])
    trm, trm_not = termin_gun(t["termin_metin"], bugun)
    gar = garanti_ay(t["garanti_metin"])
    gec = gecerlilik(t["gecerlilik_metin"], tt)
    nak = para(t["nakliye_tutar"])
    nk = katla(t["nakliye_metin"])
    if nak is None and re.search(r"\b(dahil|bize ait|tarafimiza ait|ucretsiz|bedelsiz|teslim)\b", nk):
        nak = Decimal(0)
    if v is None:
        b.append(("Dikkat", "Ödeme vadesi belgede yok ya da okunamadı; karşılaştırma kod bloğu boş vadeyi 0 gün (peşin) sayar"))
    if trm is None:
        b.append(("Dikkat", "Teslim süresi belgede yok ya da okunamadı; karşılaştırma kod bloğu boş termini 0 gün sayar, doldurun"))
    elif trm < 0:
        b.append(("Hata", f"Teslim tarihi geçmiş ({t['termin_metin']})"))
    if trm_not:
        b.append(("Bilgi", f"Termin: {trm_not}"))
    if nak is None:
        b.append(("Dikkat", f"Nakliye bedeli belirsiz ('{t['nakliye_metin'] or 'belgede yok'}'): toplam maliyete ekleyin"))
    if gar is None:
        b.append(("Bilgi", "Garanti süresi yazmıyor; karşılaştırmada 0 ay sayılır"))
    if not t["gecerlilik_metin"]:
        b.append(("Bilgi", "Teklifin geçerlilik süresi yazmıyor"))
    elif gec is None:
        b.append(("Dikkat", f"Geçerlilik ifadesi tarihe çevrilemedi: '{t['gecerlilik_metin']}'"))
    elif gec < bugun:
        b.append(("Hata", f"Teklifin geçerlilik süresi dolmuş ({gec:%d.%m.%Y}); güncel teklif isteyin"))
    elif (gec - bugun).days <= 3:
        b.append(("Dikkat", f"Teklif {gec:%d.%m.%Y} tarihinde geçerliliğini yitiriyor"))
    if t["kdv"] == "Dahil":
        b.append(("Dikkat", "Fiyatlar KDV dahil: diğer tekliflerle karşılaştırmadan önce KDV hariç tutara çevirin"))
    elif t["kdv"] == "Belirtilmemiş":
        b.append(("Dikkat", "KDV dahil mi hariç mi yazmıyor; tedarikçiye teyit ettirin"))
    if t["para_birimi"] != "TL":
        b.append(("Bilgi", f"Para birimi {t['para_birimi']}: karşılaştırmada --kur ile TL'ye çevrilir"))
    tal = {x["kod"]: x for x in talep}
    kalemler = []
    for j, k in enumerate(t["kalemler"], 1):
        m, f, tut = para(k["miktar"]), para(k["birim_fiyat"]), para(k["tutar"])
        kod = k["talep_kodu"].strip()
        uy = []
        if m is not None and f is not None and tut is not None and abs(m * f - tut) > YUVARLAMA:
            uy.append(("Hata", f"{m:g} × {tl(f)} = {tl(m * f)}, teklifte {tl(tut)}"))
        if kod in tal and m is not None and tal[kod]["miktar"] is not None and m != tal[kod]["miktar"]:
            uy.append(("Dikkat", f"Teklif miktarı {m:g}, talep {tal[kod]['miktar']:g}"))
        if m is None and kod in tal:
            m = tal[kod]["miktar"]
            uy.append(("Bilgi", "Miktar teklifte yok; talep miktarı kullanıldı"))
        if talep and not kod:
            uy.append(("Dikkat", "Talep listesindeki hiçbir malzemeyle eşleşmedi; elle kontrol edin"))
        b += [(o, f"Kalem {j} ({k['malzeme'][:30]}): {x}") for o, x in uy]
        kalemler.append(k | {"kod": kod, "talep_malzeme": tal.get(kod, {}).get("malzeme", k["malzeme"]), "m": m, "f": f, "uyarilar": uy})
    verilen = {k["kod"] for k in kalemler}
    eksik = [x for x in talep if x["kod"] not in verilen]
    if eksik:
        b.append(("Bilgi", "Teklif verilmeyen kalem: " + ", ".join(f"{x['kod']} {x['malzeme']}" for x in eksik)))
    onem = {o for o, _ in b}
    return {"vade": v, "termin": trm, "garanti": gar, "gecerlilik": gec, "nakliye": nak, "kalemler": kalemler, "bulgular": b,
            "durum": "Hatalı" if "Hata" in onem else "Kontrol" if "Dikkat" in onem else "Uygun"}


def uret(metin: str, talep: list[dict], bugun: date, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    liste = "\n".join(f"{x['kod']} | {x['malzeme']} | {x['miktar'] if x['miktar'] is not None else ''} {x['birim']}".rstrip() for x in talep)
    mesaj = (f"<talep_listesi>\n{liste or '(verilmedi: talep_kodu alanlarını boş bırakın)'}\n</talep_listesi>\n<bugun>{bugun:%d.%m.%Y}</bugun>\n"
             f"<teklif>\n{metin}\n</teklif>")
    if duzeltme:
        mesaj += ("\n<duzeltme>\nÖnceki çıkarım: " + json.dumps(onceki, ensure_ascii=False) + "\nKodun bulduğu sorunlar:\n"
                  + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")
    return llm.json_iste(sistem, mesaj, SEMA)


def duzeltme_dongusu(uret_f, denetle_f, tur: int) -> tuple[dict, list[str], int]:
    taslak = uret_f(None, None)
    kalan = denetle_f(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret_f(kalan, taslak)
        kalan = denetle_f(taslak)
    return taslak, kalan, yapilan


def isle(yol: Path, talep: list[dict], bugun: date, tur: int) -> dict:
    ham = belge.metin_oku(yol)
    if len(ham.strip()) < 20:
        raise belge.BelgeHatasi(f"{yol.name}: metin bulunamadı.")
    metin = llm.maskele(ham)
    t, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(metin, talep, bugun, d, o), lambda x: denetle(x, metin, talep), tur)
    d = {"dosya": yol.name, "teklif": t, "kalan": kalan, "tur": yapilan}
    if t["teklif_var"]:
        d.update(degerlendir(t, metin, talep, bugun))
        if kalan:
            d["durum"] = "Hatalı" if d["durum"] == "Hatalı" else "Kontrol"
    return d


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

KOD_BLOGU_SUTUNLARI = ["Tedarikçi", "Malzeme Kodu", "Malzeme", "Miktar", "Birim", "Birim Fiyat", "Para Birimi", "Nakliye", "Termin (gün)",
                       "Ödeme Vadesi (gün)", "Garanti (ay)", "Kalite Puanı", "Teknik Uygun", "Not"]


def _baslik(ws, basliklar, genislik) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font, h.alignment = BASLIK_DOLGU, BASLIK_YAZI, UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def _f(x):
    return None if x is None else float(x)


def rapor_yaz(cikti: Path, sonuc: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Teklif Listesi"
    _baslik(ws, KOD_BLOGU_SUTUNLARI, (30, 12, 34, 9, 8, 12, 8, 10, 9, 10, 9, 9, 9, 60))
    for d in sonuc:
        t = d["teklif"]
        if not t["teklif_var"]:
            continue
        for i, k in enumerate(d["kalemler"]):
            notlar = [f"Teklifteki ad: {k['malzeme']}"] if k["talep_malzeme"] != k["malzeme"] else []
            notlar += [f"{o}: {x}" for o, x in k["uyarilar"]]
            ws.append([t["tedarikci"], k["kod"], k["talep_malzeme"], _f(k["m"]), k["birim"], _f(k["f"]), t["para_birimi"],
                       _f(d["nakliye"]) if i == 0 else None, d["termin"], d["vade"], d["garanti"], None, None, "\n".join(notlar)])
            n = ws.max_row
            ws.cell(n, 6).number_format = ws.cell(n, 8).number_format = "#,##0.00"
            for j in (1, 4, 5, 6, 7):
                ws.cell(n, j).fill = AI_DOLGU
            if any(o == "Hata" for o, _ in k["uyarilar"]):
                ws.cell(n, 14).fill = KIRMIZI
            for h in ws[n]:
                h.alignment = UST

    o = wb.create_sheet("Teklifler")
    _baslik(o, ["Kaynak Dosya", "Tedarikçi", "Teklif No", "Teklif Tarihi", "Geçerlilik", "Para Birimi", "KDV", "Ödeme (belgede)", "Vade (gün)",
                "Teslim (belgede)", "Termin (gün)", "Nakliye (belgede)", "Nakliye", "Garanti (belgede)", "Garanti (ay)", "Kalem", "Özel Şartlar",
                "Durum", "Bulgular"], (26, 28, 14, 11, 11, 7, 9, 28, 7, 30, 7, 30, 10, 22, 7, 6, 40, 9, 70))
    for d in sonuc:
        t = d["teklif"]
        if not t["teklif_var"]:
            o.append([d["dosya"], "", "", "", "", "", "", "", "", "", "", "", "", "", "", 0, "", "Kontrol",
                      "Belgede fiyat teklifi bulunamadı. " + " ".join(t["notlar"])])
            continue
        notlar = [f"{x}: {m}" for x, m in d["bulgular"]] + [f"Kontrol (çıkarım): {k}" for k in d["kalan"]]
        o.append([d["dosya"], t["tedarikci"], t["teklif_no"], tarih_coz(t["teklif_tarihi"]) or t["teklif_tarihi"], d["gecerlilik"], t["para_birimi"],
                  t["kdv"], t["odeme_metin"], d["vade"], t["termin_metin"], d["termin"], t["nakliye_metin"], _f(d["nakliye"]), t["garanti_metin"],
                  d["garanti"], len(d["kalemler"]), "\n".join(t["sartlar"]), d["durum"], "\n".join(notlar)])
        n = o.max_row
        o.cell(n, 4).number_format = o.cell(n, 5).number_format = "DD.MM.YYYY"
        for j in (2, 3, 4, 7, 8, 10, 12, 14, 17):
            o.cell(n, j).fill = AI_DOLGU
        o.cell(n, 18).fill = PatternFill("solid", fgColor=RENK[d["durum"]])
        for h in o[n]:
            h.alignment = UST

    k = wb.create_sheet("Kontrol")
    _baslik(k, ["Dosya", "Düzeltme Turu", "Son kontrol"], (30, 10, 90))
    for d in sonuc:
        k.append([d["dosya"], d["tur"], "; ".join(d["kalan"]) or "Kod denetiminden sorunsuz geçti"])
        k.cell(k.max_row, 3).fill = KIRMIZI if d["kalan"] else YESIL
        k.cell(k.max_row, 3).alignment = UST
    k.append([])
    k.append([f"Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    k.append(["Önemli: Mor hücreler yapay zekâ ile çıkarılmıştır; malzeme eşleşmelerini ve fiyatları teklifin aslıyla karşılaştırın. "
              "Vade, termin, garanti ve geçerlilik günlerini kod belgedeki ifadeden hesaplamıştır. 'Teklif Listesi' sayfasında Kalite Puanı "
              "ve Teknik Uygun (E/H) sütunlarını doldurup dosyayı 'Teklif Karşılaştırma' kod bloğuna verin."])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def dosyalar(girdi: Path) -> list[Path]:
    if girdi.is_dir():
        l = sorted(p for p in girdi.iterdir() if p.is_file() and p.suffix.lower() in belge.DESTEKLENEN)
        if not l:
            raise belge.BelgeHatasi(f"{girdi}: klasörde .pdf, .docx veya .txt dosyası yok.")
        return l
    if not girdi.exists():
        raise belge.BelgeHatasi(f"{girdi}: dosya bulunamadı.")
    return [girdi]


def calistir(girdi: Path, cikti: Path, talep_yolu: Path | None = None, bugun: date = ORNEK_TARIH, tur: int = 2, evet: bool = False) -> dict:
    talep = talep_oku(talep_yolu)
    l = dosyalar(girdi)
    print(f"[OK] {len(l)} teklif belgesi · talep listesi: {len(talep)} kalem")
    llm.onay_al(f"{len(l)} teklif belgesinin metni ve talep listesi gönderilecek (telefon, e-posta, IBAN, TCKN maskeli). "
                "Firma adları, kalemler ve fiyatlar gönderilir.", evet)
    sonuc = [isle(y, talep, bugun, tur) for y in l]
    rapor_yaz(cikti, sonuc)
    js = cikti.with_suffix(".json")
    js.write_text(json.dumps(sonuc, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    return {"teklifler": sonuc, "json": js}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Tedarikçi fiyat tekliflerini (e-posta, PDF, Word) karşılaştırma tablosuna aktarır.")
    p.add_argument("--girdi", type=Path, default=ORNEK / "teklifler", help="Teklif belgesi ya da klasörü (.pdf, .docx, .txt)")
    p.add_argument("--talep", type=Path, default=None, help="Talep listesi (.xlsx/.csv: Malzeme Kodu, Malzeme, Miktar, Birim)")
    p.add_argument("--tarih", default=None, help="Termin ve geçerlilik için bugünün tarihi GG.AA.YYYY (varsayılan: bugün; örnekte 16.10.2026)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "teklifler.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    ornek = a.girdi == ORNEK / "teklifler"
    talep = a.talep or (ORNEK / "talep.csv" if ornek else None)
    bugun = tarih_coz(a.tarih) if a.tarih else (ORNEK_TARIH if ornek else date.today())
    if bugun is None:
        print(f"[X] --tarih '{a.tarih}' okunamadı (GG.AA.YYYY)")
        return 2
    try:
        s = calistir(a.girdi, a.cikti, talep, bugun, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    for d in s["teklifler"]:
        if not d["teklif"]["teklif_var"]:
            print(f"[!] {d['dosya']}: fiyat teklifi bulunamadı")
            continue
        print(f"[OK] {d['teklif']['tedarikci']}: {len(d['kalemler'])} kalem · {d['durum']}")
        for o, m in d["bulgular"]:
            if o == "Hata":
                print(f"[!]   {m}")
    print(f"[OK] Rapor: {a.cikti.resolve()} (+ {s['json'].name})")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
