"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import json
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
CEVAP = json.loads((Path(__file__).parent / "sahte_cevaplar.json").read_text(encoding="utf-8"))
KACIN = [x.strip() for x in (ORNEK / "kacinilacak_ifadeler.txt").read_text(encoding="utf-8").splitlines() if x.strip()]
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    """SW-02'de ilk turda indirimi olmayan ürüne '%20 indirim' yazar."""
    MESAJLAR.append(kullanici)
    kod = re.search(r'<urun kod="([^"]+)"', kullanici)[1]
    return copy.deepcopy(CEVAP[f"{kod}_hatali" if kod == "SW-02" and "<duzeltme>" not in kullanici else kod])


def hazir():
    return {u["kod"]: u for u in agent.urunleri_oku(ORNEK / "urunler.csv")}, agent.limit_oku(ORNEK / "kanal_limitleri.csv")


class KodTesti(unittest.TestCase):
    def test_okuma(self):
        u, lim = hazir()
        self.assertEqual(lim, {"baslik": 30, "aciklama": 90, "yol": 15})
        self.assertEqual((u["TS-01"]["indirim"], u["TS-01"]["anahtar"], u["SW-02"]["indirim"]), ("20", ["pamuklu tişört", "kadın tişört"], ""))

    def test_denetim(self):
        u, lim = hazir()
        d = lambda k, c, nb=10: agent.denetle(u[k], c, lim, nb, 4, KACIN, 3)
        self.assertEqual([d("TS-01", CEVAP["TS-01"]), d("SW-02", CEVAP["SW-02"])], [[], []])
        self.assertEqual(d("SW-02", CEVAP["SW-02_hatali"]),
                         ["SW-02: ürün bilgisinde olmayan sayı: 20", "SW-02: indirim / kampanya alanı boş ama metinde 'indirimle' var"])
        c = copy.deepcopy(CEVAP["TS-01"])
        c["basliklar"][2] = "Basic Pamuklu Tişört Şimdi Burada"           # 33 karakter
        c["basliklar"][1] = "basic pamuklu tişört"                         # tekrar
        c["basliklar"][9] = "EN İYİ TİŞÖRT"                                # kaçınılacak + büyük harf
        c["aciklamalar"].pop()
        c["yol1"] = "kadın tişört"
        s = d("TS-01", c)
        for p in ("4 açıklama istendi", "başlık 3 33 karakter", "tekrar eden başlık", "kaçınılacak ifade: 'en iyi'", "tamamı büyük harf", "yol 1 'kadın tişört'"):
            self.assertTrue(any(p in x for x in s), p)
        c = copy.deepcopy(CEVAP["SW-02"])
        c["basliklar"] = ["Sıcak Tutan Model"] * 1 + [f"Model {i}" for i in range(2, 11)]
        self.assertTrue(any("anahtar kelime yalnız 0 başlıkta" in x for x in d("SW-02", c)))
        self.assertTrue(any("olmayan sayı" in x for x in d("SW-02", c)))           # "Model 2" … ürün bilgisinde yok

    def test_yuzde_100_istisnasi(self):
        u, lim = hazir()
        self.assertIn("%100 pamuk", agent.bilgi(u["TS-01"]))
        self.assertNotIn("kaçınılacak", " ".join(agent.denetle(u["TS-01"], CEVAP["TS-01"], lim, 10, 4, KACIN, 3)))
        c = copy.deepcopy(CEVAP["SW-02"])
        c["aciklamalar"][0] = "%100 memnuniyet garantili kapüşonlu sweatshirt."
        s = agent.denetle(u["SW-02"], c, lim, 10, 4, KACIN, 3)
        self.assertTrue(any("'garantili'" in x for x in s) and any("'%100'" in x for x in s))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ornek_ve_duzeltme(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "r.xlsx"
            self.assertEqual(agent.main(["--evet", "--cikti", str(cikti)]), 0)
            self.assertEqual(len(MESAJLAR), 3)
            self.assertIn("10 başlık (en fazla 30 karakter)", MESAJLAR[0])
            wb = load_workbook(cikti)
            ws = wb["Reklamlar"]
            self.assertEqual([c.value for c in ws[1]][:3] + [ws.cell(1, 13).value, ws.cell(1, 17).value], ["Ürün Kodu", "Ürün Adı", "Headline 1", "Description 1", "Path 1"])
            self.assertEqual(ws.cell(3, 19).value, "Geçti")
            self.assertNotIn("%20", " ".join(str(c.value) for c in ws[3]))
            self.assertEqual(wb["Karakter Sayıları"].max_row, 1 + 2 * 16)

    def test_cli(self):
        self.assertEqual(agent.main(["--baslik", "20", "--evet"]), 1)
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(agent.main(["--urunler", str(Path(tmp) / "yok.csv"), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
