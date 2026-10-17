"""
Reklam Metni Varyasyonları — Gerisi Bende AI Agent
E-Ticaret › Performans Pazarlama Uzmanı

Ürün bilgisinden arama ağı reklamları (Google Ads duyarlı arama reklamı) için başlık, açıklama ve görünen URL yolu
varyasyonları yazar; karakter sınırlarını ve iddiaları kod denetler.
1. Kod ürün dosyasını (Ürün Kodu; Ürün Adı; Kategori; Özellikler; Fiyat; İndirim (%); Kampanya; Anahtar Kelimeler),
   kanal sınırlarını (Alan; Azami Karakter) ve kaçınılacak ifadeleri okur.
2. Model her ürün için istenen sayıda başlık ve açıklama ile iki yol alanı yazar.
3. Kod denetler: adet; karakter sınırı; tekrar eden başlık; ürün bilgisinde olmayan sayı (fiyat, indirim oranı …);
   indirim / kampanya alanı boşken indirim dili; kaçınılacak ifadeler (ürün bilgisinde geçmiyorsa); yeterli sayıda
   başlıkta anahtar kelime; tamamı büyük harf; yol biçimi. Sorun varsa yalnız o ürün modele geri gönderilir.
4. Kod karakter sayılarını yazar ve Google Ads Editor'e kopyalanabilir sütun düzeninde tablo üretir.

Kullanım:
    python agent.py                                                     # örnek 2 ürün
    python agent.py --urunler urunler.xlsx --baslik 15 --aciklama 4
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import llm

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}
INDIRIM = re.compile(r"\b(indirim\w*|kampanya\w*|firsat\w*|sezon sonu|ucuzluk|outlet|kelepir)\b")
VARSAYILAN_LIMIT = {"baslik": 30, "aciklama": 90, "yol": 15}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def sayi(x) -> Decimal | None:
    s = re.sub(r"[^\d,.\-]", "", str(x if x is not None else ""))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1 or re.fullmatch(r"\d{1,3}\.\d{3}", s):
        s = s.replace(".", "")
    try:
        return Decimal(s) if s not in ("", "-", ".") else None
    except InvalidOperation:
        return None


def sayilar(metin: str) -> set[Decimal]:
    return {v for x in re.findall(r"\d+(?:[.,]\d+)*", str(metin)) if (v := sayi(x)) is not None}


# ----------------------------------------------------------------------------
# Girdi
# ----------------------------------------------------------------------------

def tablo_oku(yol: Path) -> list[dict]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, read_only=True, data_only=True)
        s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=metin.split("\n")[0].count)))
    s = [r for r in s if any(c not in (None, "") for c in r)]
    bas = [katla(c) for c in s[0]] if s else []
    return [dict(zip(bas, list(r) + [None] * (len(bas) - len(r)))) for r in s[1:]]


def m(x) -> str:
    return "" if x is None else str(x).strip()


def urunleri_oku(yol: Path) -> list[dict]:
    out = []
    for r in tablo_oku(yol):
        kod, ad = m(r.get("urun kodu") or r.get("kod") or r.get("sku")), m(r.get("urun adi") or r.get("ad"))
        if not kod or not ad:
            continue
        out.append({"kod": kod, "ad": ad, "kategori": m(r.get("kategori")), "ozellikler": m(r.get("ozellikler") or r.get("ozellik")),
                    "fiyat": m(r.get("fiyat tl") or r.get("fiyat")), "indirim": m(r.get("indirim %") or r.get("indirim")), "kampanya": m(r.get("kampanya")),
                    "anahtar": [x.strip() for x in m(r.get("anahtar kelimeler")).split(",") if x.strip()]})
    if not out:
        raise llm.LLMHatasi(f"{yol.name}: 'Ürün Kodu' ve 'Ürün Adı' sütunlarıyla en az bir ürün gerekli.")
    if len({u["kod"] for u in out}) != len(out):
        raise llm.LLMHatasi(f"{yol.name}: ürün kodu tekrar ediyor.")
    return out


def limit_oku(yol: Path | None) -> dict[str, int]:
    lim = dict(VARSAYILAN_LIMIT)
    if yol:
        for r in tablo_oku(yol):
            a, v = katla(r.get("alan")), sayi(r.get("azami karakter"))
            k = {"baslik": "baslik", "aciklama": "aciklama", "yol": "yol"}.get(a)
            if k and v:
                lim[k] = int(v)
    return lim


def bilgi(u: dict) -> str:
    return " ".join([u["ad"], u["kategori"], u["ozellikler"], u["fiyat"], f"%{u['indirim']}" if u["indirim"] else "", u["kampanya"]])


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

SEMA = {"type": "object", "properties": {"basliklar": {"type": "array", "items": S}, "aciklamalar": {"type": "array", "items": S}, "yol1": S, "yol2": S},
        "required": ["basliklar", "aciklamalar", "yol1", "yol2"], "additionalProperties": False}


def denetle(u: dict, c: dict, lim: dict, n_b: int, n_a: int, kacin: list[str], anahtar_min: int) -> list[str]:
    k, sorun = u["kod"], []
    if len(c["basliklar"]) != n_b or len(c["aciklamalar"]) != n_a:
        sorun.append(f"{k}: {n_b} başlık ve {n_a} açıklama istendi; {len(c['basliklar'])} başlık, {len(c['aciklamalar'])} açıklama var")
    for ad, liste, sinir in (("başlık", c["basliklar"], lim["baslik"]), ("açıklama", c["aciklamalar"], lim["aciklama"]), ("yol", [c["yol1"], c["yol2"]], lim["yol"])):
        for i, t in enumerate(liste, 1):
            if len(t) > sinir:
                sorun.append(f"{k}: {ad} {i} {len(t)} karakter (sınır {sinir}): \"{t}\"")
            if not t.strip() and ad != "yol":
                sorun.append(f"{k}: {ad} {i} boş")
    kb = [katla(t) for t in c["basliklar"]]
    tekrar = sorted({t for t in c["basliklar"] if kb.count(katla(t)) > 1})
    if tekrar:
        sorun.append(f"{k}: tekrar eden başlık: {', '.join(tekrar)}")
    hepsi = " ".join(c["basliklar"] + c["aciklamalar"])
    fazla = sorted(sayilar(hepsi) - sayilar(bilgi(u)), key=float)
    if fazla:
        sorun.append(f"{k}: ürün bilgisinde olmayan sayı: {', '.join(f'{x:g}' for x in fazla)}")
    if not (u["indirim"] or u["kampanya"]) and (g := INDIRIM.search(katla(hepsi))):
        sorun.append(f"{k}: indirim / kampanya alanı boş ama metinde '{g[0]}' var")
    kbi = katla(bilgi(u))
    for x in kacin:
        kx = katla(x)
        if kx and re.search(rf"(?<![a-z0-9]){re.escape(kx)}(?![a-z0-9])", katla(hepsi)) and kx not in kbi:
            sorun.append(f"{k}: kaçınılacak ifade: '{x}'")
    if u["anahtar"]:
        n = sum(any(katla(a) in t for a in u["anahtar"]) for t in kb)
        if n < min(anahtar_min, n_b):
            sorun.append(f"{k}: anahtar kelime yalnız {n} başlıkta var; en az {min(anahtar_min, n_b)} başlıkta kullanın ({', '.join(u['anahtar'])})")
    buyuk = [t for t in c["basliklar"] + c["aciklamalar"] if len(re.sub(r"[^A-Za-zÇĞİÖŞÜçğıöşü]", "", t)) > 5 and t == t.upper()]
    if buyuk:
        sorun.append(f"{k}: tamamı büyük harf metin: {', '.join(buyuk)}")
    for i, y in enumerate((c["yol1"], c["yol2"]), 1):
        if y and not re.fullmatch(r"[a-z0-9-]+", y):
            sorun.append(f"{k}: yol {i} '{y}' yalnız küçük harf, rakam ve tire içermeli (Türkçe karakter ve boşluk yok)")
    return sorun


def uret(urunler: list[dict], lim: dict, n_b: int, n_a: int, kacin: list[str], duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    sonuc = dict(onceki or {})
    hedef = urunler if duzeltme is None else [u for u in urunler if any(s.startswith(u["kod"] + ":") for s in duzeltme)]
    for u in hedef:
        mesaj = (f"<urun kod=\"{u['kod']}\">\nAd: {u['ad']}\nKategori: {u['kategori']}\nÖzellikler: {u['ozellikler']}\nFiyat: {u['fiyat'] or '-'} TL\n"
                 f"İndirim: {('%' + u['indirim']) if u['indirim'] else '-'}\nKampanya: {u['kampanya'] or '-'}\nAnahtar kelimeler: {', '.join(u['anahtar']) or '-'}\n</urun>\n"
                 f"<istek>{n_b} başlık (en fazla {lim['baslik']} karakter), {n_a} açıklama (en fazla {lim['aciklama']} karakter), "
                 f"2 yol (en fazla {lim['yol']} karakter)</istek>\n<kacinilacak>{', '.join(kacin)}</kacinilacak>")
        d = [s.split(": ", 1)[1] for s in (duzeltme or []) if s.startswith(u["kod"] + ":")]
        if d:
            mesaj += ("\n<duzeltme>\nÖnceki cevap: " + json.dumps(sonuc.get(u["kod"]), ensure_ascii=False) + "\nKodun bulduğu sorunlar:\n"
                      + "\n".join(f"- {s}" for s in d) + "\n</duzeltme>")
        sonuc[u["kod"]] = llm.json_iste(sistem, mesaj, SEMA)
    return sonuc


def duzeltme_dongusu(uret_f, denetle_f, tur: int) -> tuple[dict, list[str], int]:
    taslak = uret_f(None, None)
    kalan = denetle_f(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret_f(kalan, taslak)
        kalan = denetle_f(taslak)
    return taslak, kalan, yapilan


# ----------------------------------------------------------------------------
# Çıktı
# ----------------------------------------------------------------------------

def rapor_yaz(cikti: Path, urunler: list[dict], sonuc: dict, lim: dict, n_b: int, n_a: int, kalan: list[str], tur: int) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Reklamlar"
    bas = ["Ürün Kodu", "Ürün Adı"] + [f"Headline {i}" for i in range(1, n_b + 1)] + [f"Description {i}" for i in range(1, n_a + 1)] + ["Path 1", "Path 2", "Son Kontrol"]
    ws.append(bas)
    for h in ws[1]:
        h.fill, h.font, h.alignment = BASLIK_DOLGU, BASLIK_YAZI, UST
    for u in urunler:
        c = sonuc[u["kod"]]
        b = (c["basliklar"] + [""] * n_b)[:n_b]
        a = (c["aciklamalar"] + [""] * n_a)[:n_a]
        sor = "; ".join(x.split(": ", 1)[1] for x in kalan if x.startswith(u["kod"] + ":"))
        ws.append([u["kod"], u["ad"], *b, *a, c["yol1"], c["yol2"], sor or "Geçti"])
        r = ws.max_row
        for j, (t, s) in enumerate([(x, lim["baslik"]) for x in b] + [(x, lim["aciklama"]) for x in a] + [(c["yol1"], lim["yol"]), (c["yol2"], lim["yol"])], 3):
            ws.cell(r, j).fill = KIRMIZI if len(t) > s else AI_DOLGU
        ws.cell(r, len(bas)).fill = YESIL if not sor else KIRMIZI
    for j in range(1, len(bas) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 16 if j <= 2 else (36 if j > 2 + n_b and j <= 2 + n_b + n_a else 26)
    ws.freeze_panes = "C2"
    g = wb.create_sheet("Karakter Sayıları")
    g.append(["Ürün Kodu", "Tür", "No", "Metin", "Karakter", "Sınır", "Kalan"])
    for h in g[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for u in urunler:
        c = sonuc[u["kod"]]
        for tur_ad, liste, s in (("Başlık", c["basliklar"], lim["baslik"]), ("Açıklama", c["aciklamalar"], lim["aciklama"]), ("Yol", [c["yol1"], c["yol2"]], lim["yol"])):
            for i, t in enumerate(liste, 1):
                g.append([u["kod"], tur_ad, i, t, len(t), s, s - len(t)])
                g.cell(g.max_row, 7).fill = KIRMIZI if len(t) > s else YESIL
    g.column_dimensions["D"].width = 90
    k = wb.create_sheet("Kontrol")
    k.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    k.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    k.cell(2, 1).fill = KIRMIZI if kalan else YESIL
    k.append([f"Sınırlar: başlık {lim['baslik']}, açıklama {lim['aciklama']}, yol {lim['yol']} karakter (kanal limitleri dosyası). Mor hücreler modelin "
              "metinleridir. Yayınlamadan önce reklam politikası, marka dili ve ürün bilgisinin güncelliğini kontrol edin; fiyat ve indirim "
              "değiştiyse metinleri yeniden üretin. Google Ads başlıkları herhangi bir sırayla birleştirebilir."])
    k.column_dimensions["A"].width = 140
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(u_yol: Path, cikti: Path, l_yol: Path | None, kacin: list[str], n_b: int = 10, n_a: int = 4, anahtar_min: int = 3, tur: int = 2, evet: bool = False) -> dict:
    if not 3 <= n_b <= 15 or not 2 <= n_a <= 4:
        raise llm.LLMHatasi("Başlık sayısı 3–15, açıklama sayısı 2–4 olmalı (duyarlı arama reklamı).")
    urunler, lim = urunleri_oku(u_yol), limit_oku(l_yol)
    print(f"[OK] {len(urunler)} ürün · sınırlar: başlık {lim['baslik']}, açıklama {lim['aciklama']}, yol {lim['yol']}")
    llm.onay_al(f"{len(urunler)} ürünün adı, özellikleri, fiyatı ve kampanya bilgisi gönderilecek.", evet)
    sonuc, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(urunler, lim, n_b, n_a, kacin, d, o),
                                             lambda x: [s for u in urunler for s in denetle(u, x[u["kod"]], lim, n_b, n_a, kacin, anahtar_min)], tur)
    rapor_yaz(cikti, urunler, sonuc, lim, n_b, n_a, kalan, yapilan)
    return {"urunler": urunler, "sonuc": sonuc, "kalan": kalan, "tur": yapilan, "limit": lim}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Ürün bilgisinden arama ağı reklam başlık / açıklama varyasyonları yazar; sınır ve iddiaları kod denetler.")
    p.add_argument("--urunler", type=Path, default=ORNEK / "urunler.csv", help="Ürün Kodu; Ürün Adı; Kategori; Özellikler; Fiyat; İndirim (%%); Kampanya; Anahtar Kelimeler")
    p.add_argument("--limitler", type=Path, default=ORNEK / "kanal_limitleri.csv", help="Alan; Azami Karakter (Başlık / Açıklama / Yol)")
    p.add_argument("--kacinilacak", type=Path, default=ORNEK / "kacinilacak_ifadeler.txt", help="Kaçınılacak ifadeler, satır başına bir tane")
    p.add_argument("--baslik", type=int, default=10, help="Ürün başına başlık sayısı, 3–15 (varsayılan 10)")
    p.add_argument("--aciklama", type=int, default=4, help="Ürün başına açıklama sayısı, 2–4 (varsayılan 4)")
    p.add_argument("--anahtar-min", type=int, default=3, help="Anahtar kelime geçmesi gereken en az başlık sayısı (varsayılan 3)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "reklam_metinleri.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        if not a.kacinilacak.exists():
            raise llm.LLMHatasi(f"{a.kacinilacak}: kaçınılacak ifadeler dosyası bulunamadı.")
        kacin = [x.strip() for x in a.kacinilacak.read_text(encoding="utf-8-sig").splitlines() if x.strip()]
        s = calistir(a.urunler, a.cikti, a.limitler, kacin, a.baslik, a.aciklama, a.anahtar_min, a.tur, a.evet)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    for u in s["urunler"]:
        c = s["sonuc"][u["kod"]]
        print(f"[OK] {u['kod']} · {len(c['basliklar'])} başlık (en uzun {max(map(len, c['basliklar']), default=0)}) · {len(c['aciklamalar'])} açıklama")
    for x in s["kalan"]:
        print(f"[!] Kalan: {x}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
