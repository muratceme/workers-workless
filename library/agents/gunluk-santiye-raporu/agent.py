"""
Günlük Şantiye Raporu — Gerisi Bende AI Agent
Şantiye › Saha Mühendisi

1. Kod saha notlarını satır satır numaralar; İSG anahtar kelimesi geçen satırları bulur (kaza, yaralanma, burkulma,
   düşme, çarpma, korkuluk, iskele, emniyet kemeri, baret, yangın, elektrik çarpması, göçük, ramak kala).
2. Model günlük raporu bölümlere ayırır: genel bilgi ve hava durumu, personel (ekip / taşeron ve sayı), makine-ekipman,
   yapılan imalatlar (yer, iş kalemi, miktar, birim), malzeme girişleri, sorunlar ve gecikmeler, İSG olayları,
   ziyaretçi / denetimler, ertesi günün planı. Her madde nottaki satır numaralarına dayanır.
3. Kod denetler: her madde geçerli satıra dayanıyor mu; kişi, makine, miktar ve sıcaklık değerleri dayandığı satırda
   yazıyor mu; İSG satırlarının hepsi İSG bölümünde mi; hiçbir bölüme girmeyen satır kalmış mı. Sorun varsa modele geri
   gönderilir (en fazla --tur kez). Toplam personel ve makine sayısını kod hesaplar.

Kullanım:
    python agent.py                                         # örnek saha notlarıyla dener
    python agent.py --girdi notlar.docx --tur 3
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
ISG = re.compile(r"\b(kaza|yaralan|burk|dustu|dusme|carp|korkuluk|iskele|emniyet kemeri|baret|yangin|elektrik carp|gocuk|ramak kala|kirildi|kesildi)")
K = {"type": "array", "items": {"type": "integer"}}
N = {"type": ["number", "null"]}


def _m(**alan) -> dict:
    alan["kaynak_satirlar"] = K
    return {"type": "array", "items": {"type": "object", "properties": alan, "required": list(alan), "additionalProperties": False}}


SEMA = {
    "type": "object",
    "properties": {
        "genel": {"type": "object", "properties": {"tarih": S, "proje": S, "hazirlayan": S, "kaynak_satirlar": K},
                  "required": ["tarih", "proje", "hazirlayan", "kaynak_satirlar"], "additionalProperties": False},
        "hava": {"type": "object", "properties": {"durum": S, "sicaklik_c": N, "calismaya_etkisi": S, "kaynak_satirlar": K},
                 "required": ["durum", "sicaklik_c", "calismaya_etkisi", "kaynak_satirlar"], "additionalProperties": False},
        "personel": _m(ekip=S, sayi=N, is_yeri=S),
        "makineler": _m(makine=S, sayi=N, aciklama=S),
        "imalatlar": _m(yer=S, is_kalemi=S, miktar=N, birim=S, durum=S),
        "malzeme": _m(malzeme=S, miktar=N, birim=S, aciklama=S),
        "sorunlar": _m(konu=S, etki=S, aksiyon=S),
        "isg": _m(olay=S, yapilan=S, takip=S),
        "ziyaretler": _m(kim=S, amac=S, sonuc=S),
        "yarin": _m(is_tanimi=S),
    },
    "required": ["genel", "hava", "personel", "makineler", "imalatlar", "malzeme", "sorunlar", "isg", "ziyaretler", "yarin"],
    "additionalProperties": False,
}
BOLUMLER = (("personel", "Personel"), ("makineler", "Makine-ekipman"), ("imalatlar", "Yapılan imalatlar"), ("malzeme", "Malzeme girişleri"),
            ("sorunlar", "Sorunlar ve gecikmeler"), ("isg", "İSG"), ("ziyaretler", "Ziyaret / denetim"), ("yarin", "Yarının planı"))
SAYI_ALANI = {"personel": "sayi", "makineler": "sayi", "imalatlar": "miktar", "malzeme": "miktar"}


def satirlar(metin: str) -> list[str]:
    return [s.strip() for s in metin.splitlines() if s.strip()]


def isg_satirlari(l: list[str]) -> list[int]:
    return [i for i, s in enumerate(l, 1) if ISG.search(katla(s))]


def _ilk(m: dict) -> str:
    return next((v for k, v in m.items() if k != "kaynak_satirlar" and isinstance(v, str) and v), "")


def denetle(t: dict, l: list[str]) -> list[str]:
    n, sorun, kullanilan = len(l), [], set()

    def kaynak(ad: str, ks: list[int]) -> list[str]:
        if not ks:
            sorun.append(f"{ad}: kaynak satır numarası yok")
        gec = [k for k in ks if not 1 <= k <= n]
        if gec:
            sorun.append(f"{ad}: notlarda olmayan satır {gec}")
        kullanilan.update(ks)
        return [l[k - 1] for k in ks if 1 <= k <= n]

    kaynak("Genel bilgi", t["genel"]["kaynak_satirlar"])
    hs = kaynak("Hava", t["hava"]["kaynak_satirlar"])
    if t["hava"]["sicaklik_c"] is not None and sayi_yaz(t["hava"]["sicaklik_c"]) not in set().union(*map(sayilar, hs or [""])):
        sorun.append(f"Hava: sıcaklık {sayi_yaz(t['hava']['sicaklik_c'])} dayandığı satırda yazmıyor")
    for anahtar, ad in BOLUMLER:
        for m in t[anahtar]:
            etiket = f"{ad} · '{_ilk(m)[:40]}'"
            sat = kaynak(etiket, m["kaynak_satirlar"])
            alan = SAYI_ALANI.get(anahtar)
            if alan and m[alan] is not None and sayi_yaz(m[alan]) not in set().union(*map(sayilar, sat or [""])):
                sorun.append(f"{etiket}: {alan} {sayi_yaz(m[alan])} dayandığı satırda yazmıyor (hesaplamayın, tahmin etmeyin)")
    isg_k = {k for m in t["isg"] for k in m["kaynak_satirlar"]}
    for i in isg_satirlari(l):
        if i not in isg_k:
            sorun.append(f"Satır {i} İSG ile ilgili ('{l[i - 1][:70]}') ama İSG bölümünde yok")
    for i, s in enumerate(l, 1):
        if i not in kullanilan and len(s.split()) >= 3:
            sorun.append(f"Satır {i} hiçbir bölümde kullanılmadı: '{s[:70]}'")
    return sorun


def uret(l: list[str], duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    numarali = "\n".join(f"{i}: {llm.maskele(s)}" for i, s in enumerate(l, 1))
    bulgu = "İSG anahtar kelimesi geçen satırlar: " + (", ".join(map(str, isg_satirlari(l))) or "yok")
    return llm.json_iste(sistem, f"<saha_notlari>\n{numarali}\n</saha_notlari>\n<kod_bulgulari>{bulgu}</kod_bulgulari>" + duzeltme_blogu(duzeltme, onceki), SEMA)


def toplam(liste: list[dict]) -> float:
    return sum(m["sayi"] or 0 for m in liste)


def rapor_metni(t: dict) -> str:
    g, h = t["genel"], t["hava"]
    out = [f"# Günlük Şantiye Raporu — {g['proje']} · {g['tarih']}", f"Hazırlayan: {g['hazirlayan']}", "",
           f"**Hava:** {h['durum']}" + (f", {sayi_yaz(h['sicaklik_c'])} °C" if h["sicaklik_c"] is not None else "") + (f" — {h['calismaya_etkisi']}" if h["calismaya_etkisi"] else ""),
           f"**Toplam personel:** {sayi_yaz(toplam(t['personel']))} · **Toplam makine:** {sayi_yaz(toplam(t['makineler']))}"]
    for anahtar, ad in BOLUMLER:
        if t[anahtar]:
            out += ["", f"## {ad}"]
            for m in t[anahtar]:
                parca = [str(v) if not isinstance(v, float) else sayi_yaz(v) for k, v in m.items() if k != "kaynak_satirlar" and v not in ("", None)]
                out.append(f"- {' · '.join(parca)} _(satır {', '.join(map(str, m['kaynak_satirlar']))})_")
    return "\n".join(out)


def rapor_yaz(cikti: Path, l, t, kalan, tur) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Günlük Rapor"
    g, h = t["genel"], t["hava"]
    ws.append([f"Günlük Şantiye Raporu · {g['proje']} · {g['tarih']} · Hazırlayan: {g['hazirlayan']}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Hava: {h['durum']}" + (f", {sayi_yaz(h['sicaklik_c'])} °C" if h["sicaklik_c"] is not None else "") + (f" · {h['calismaya_etkisi']}" if h["calismaya_etkisi"] else "")])
    ws.append([f"Toplam personel: {sayi_yaz(toplam(t['personel']))} · Toplam makine: {sayi_yaz(toplam(t['makineler']))}"])
    ws.append([])
    baslik(ws, ["Bölüm", "Konu", "Sayı / Miktar", "Birim", "Detay", "Not Satırı"], (22, 34, 12, 8, 60, 10))
    for anahtar, ad in BOLUMLER:
        for m in t[anahtar]:
            deger = m.get(SAYI_ALANI.get(anahtar, ""), None)
            konu = m.get("ekip") or m.get("makine") or (f"{m['yer']} · {m['is_kalemi']}" if "yer" in m else None) or m.get("malzeme") or m.get("konu") or m.get("olay") or m.get("kim") or m.get("is_tanimi")
            detay = " · ".join(str(v) for k, v in m.items() if k not in ("kaynak_satirlar", "ekip", "makine", "yer", "is_kalemi", "malzeme", "konu", "olay", "kim", "is_tanimi",
                                                                         "sayi", "miktar", "birim") and v)
            ws.append([ad, konu, deger, m.get("birim", "adet" if anahtar == "makineler" else ("kişi" if anahtar == "personel" else "")), detay,
                       ", ".join(map(str, m["kaynak_satirlar"]))])
            if anahtar in ("isg", "sorunlar"):
                boya(ws, [1], KIRMIZI if anahtar == "isg" else SARI)
    son_not(ws, "Önemli: Yapay zekâ taslağıdır. Şantiye şefi notlarla karşılaştırıp onaylamalıdır; resmî şantiye defteri ve iş kazası bildirimi ayrıca yapılır.",
            kalan, tur)
    hizala(ws, 5)
    d = wb.create_sheet("Notlar")
    baslik(d, ["Satır", "Not", "Kullanıldığı Bölüm", "İSG"], (6, 90, 30, 6))
    yer = {}
    for k in t["genel"]["kaynak_satirlar"]:
        yer.setdefault(k, []).append("Genel")
    for k in t["hava"]["kaynak_satirlar"]:
        yer.setdefault(k, []).append("Hava")
    for anahtar, ad in BOLUMLER:
        for m in t[anahtar]:
            for k in m["kaynak_satirlar"]:
                yer.setdefault(k, []).append(ad)
    isg = set(isg_satirlari(l))
    for i, s in enumerate(l, 1):
        d.append([i, s, ", ".join(dict.fromkeys(yer.get(i, []))) or "—", "Evet" if i in isg else ""])
        if not yer.get(i) and len(s.split()) >= 3:
            boya(d, [3], KIRMIZI)
        if i in isg:
            boya(d, [4], SARI)
    hizala(d)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    md = cikti.with_suffix(".md")
    md.write_text(rapor_metni(t) + "\n", encoding="utf-8")
    return md


def calistir(girdi: Path, cikti: Path, tur: int = 2, evet: bool = False) -> dict:
    if not girdi.exists():
        raise belge.BelgeHatasi(f"{girdi}: dosya bulunamadı.")
    l = satirlar(belge.metin_oku(girdi))
    if not l:
        raise belge.BelgeHatasi(f"{girdi.name}: metin bulunamadı.")
    print(f"[OK] {len(l)} satır · İSG satırı {len(isg_satirlari(l))}")
    llm.onay_al("Saha notları (numaralı satırlar; telefon / e-posta maskeli) rapor yazımı için gönderilecek.", evet)
    t, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(l, d, o), lambda x: denetle(x, l), tur)
    md = rapor_yaz(cikti, l, t, kalan, yapilan)
    return {"satirlar": l, "taslak": t, "kalan": kalan, "tur": yapilan, "md": md}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    p = argparse.ArgumentParser(description="Saha notlarından hava, personel-makine, imalat, sorun ve İSG bölümlü günlük şantiye raporu hazırlar.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "saha_notlari.txt", help="Saha notları (.txt, .md, .docx, .pdf)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "gunluk_santiye_raporu.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.cikti, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    t = s["taslak"]
    print(f"[OK] Rapor hazır · personel {sayi_yaz(toplam(t['personel']))} · imalat {len(t['imalatlar'])} · sorun {len(t['sorunlar'])} · İSG {len(t['isg'])}")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()} (+ {s['md'].name})")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
