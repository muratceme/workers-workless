"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri" / "saha_notlari.txt"
MESAJLAR = []


def taslak(dogru=True) -> dict:
    t = {"genel": {"tarih": "15.10.2026", "proje": "Örnek Konut Projesi B Blok", "hazirlayan": "K. Aydın", "kaynak_satirlar": [1]},
         "hava": {"durum": "Sabah sisli, öğleden sonra açık", "sicaklik_c": 14, "calismaya_etkisi": "", "kaynak_satirlar": [2]},
         "personel": [{"ekip": "Kaba inşaat taşeronu", "sayi": 18, "is_yeri": "B Blok", "kaynak_satirlar": [3]},
                      {"ekip": "Demirci", "sayi": 9, "is_yeri": "B Blok", "kaynak_satirlar": [3]},
                      {"ekip": "Kalıpçı", "sayi": 12, "is_yeri": "B Blok", "kaynak_satirlar": [3]},
                      {"ekip": "Elektrik taşeronu", "sayi": 4, "is_yeri": "A Blok tesisat", "kaynak_satirlar": [4]}],
         "makineler": [{"makine": "Kule vinç", "sayi": 1, "aciklama": "", "kaynak_satirlar": [5]},
                       {"makine": "Beton pompası", "sayi": 1, "aciklama": "13:00-17:30", "kaynak_satirlar": [5]}],
         "imalatlar": [{"yer": "B Blok 6. kat", "is_kalemi": "Döşeme betonu C30/37", "miktar": 86, "birim": "m3", "durum": "Tamamlandı, 6 küp numune", "kaynak_satirlar": [6]},
                       {"yer": "B Blok 7. kat", "is_kalemi": "Kolon demir montajı", "miktar": 40, "birim": "%", "durum": "Devam ediyor", "kaynak_satirlar": [7]}],
         "malzeme": [{"malzeme": "Ø12 nervürlü demir", "miktar": 18, "birim": "ton", "aciklama": "İrsaliye kontrol edildi", "kaynak_satirlar": [9]}],
         "sorunlar": [{"konu": "Hazır beton mikseri gecikmesi", "etki": "Döküm 45 dk aksadı", "aksiyon": "Beton firmasına bildirilmeli", "kaynak_satirlar": [8]}],
         "isg": [{"olay": "Kalıpçı işçi iskeleden inerken ayağını burktu", "yapilan": "Revirde müdahale, iş kazası formu", "takip": "", "kaynak_satirlar": [10]},
                 {"olay": "6. kat döşeme kenarında korkuluk eksik", "yapilan": "", "takip": "Yarına tamamlanacak", "kaynak_satirlar": [11]}],
         "ziyaretler": [{"kim": "Yapı denetim kontrol mühendisi", "amac": "Demir kontrolü", "sonuc": "7. kat kolon demirleri onaylandı", "kaynak_satirlar": [12]}],
         "yarin": [{"is_tanimi": "7. kat kolon betonu; pompa rezervasyonu teyit edilmeli", "kaynak_satirlar": [13]}]}
    if not dogru:
        t["isg"].pop()
        t["sorunlar"].append({"konu": "Korkuluk eksik", "etki": "", "aksiyon": "Tamamlanacak", "kaynak_satirlar": [11]})
        t["personel"][2]["sayi"] = 14
        t["ziyaretler"] = []
    return t


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return taslak("<duzeltme>" in kullanici)


class KodTesti(unittest.TestCase):
    def test_denetim(self):
        l = agent.satirlar(ORNEK.read_text(encoding="utf-8"))
        self.assertEqual((len(l), agent.isg_satirlari(l)), (13, [10, 11]))
        self.assertEqual(agent.denetle(taslak(), l), [])
        s = agent.denetle(taslak(False), l)
        self.assertEqual(len(s), 3)
        self.assertTrue(any(x.startswith("Satır 11 İSG") for x in s) and any("sayi 14" in x for x in s) and any(x.startswith("Satır 12 hiçbir") for x in s))
        t = taslak()
        t["hava"]["sicaklik_c"] = 16
        t["imalatlar"][0]["kaynak_satirlar"] = [6, 30]
        self.assertEqual(len(agent.denetle(t, l)), 2)
        self.assertEqual((agent.toplam(taslak()["personel"]), agent.toplam(taslak()["makineler"])), (43, 2))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "s.xlsx"
            s = agent.calistir(ORNEK, cikti, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            self.assertIn("İSG anahtar kelimesi geçen satırlar: 10, 11", MESAJLAR[0])
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Günlük Rapor", "Notlar"])
            self.assertEqual(wb["Günlük Rapor"]["A3"].value, "Toplam personel: 43 · Toplam makine: 2")
            md = s["md"].read_text(encoding="utf-8")
            self.assertIn("## İSG", md)
            self.assertIn("14 °C", md)

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(ORNEK, Path(tmp) / "s.xlsx", tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 3)
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--girdi", str(ORNEK.with_name("yok.txt")), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
