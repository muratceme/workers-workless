"""
Reasürans Slipi Özeti ve Kontrolü — Workers / Workless AI Agent
Reasürans › Reasürans Uzmanı (Reasürans Müdürü de kullanır)

Reasürans slipi / anlaşma metninden ana şartları özet tabloya çıkarır, beklenen (onaylanan) şartlarla karşılaştırır ve
zorunlu klozların varlığını kontrol eder.
1. Kod slip metnini, beklenen şartlar tablosunu ve zorunlu kloz listesini okur; telefon, e-posta, IBAN maskelenir.
2. Model ana şartları belgedeki yazımıyla çıkarır: tür, dönem, kapsam, limit, öncelik, para birimi, yeniden tesis, MDP,
   prim oranı, saat klozu, istisnalar, hasar bildirimi, uyuşmazlık, reasürörler ve payları.
3. Kod çıkarımı denetler: tarih ve sayılar slipte geçiyor mu; reasürör adları slipte aynen var mı. Sorun varsa modele geri
   gönderilir (en fazla --tur kez).
4. Kod kontrolleri yapar: reasürör payları toplamı %100 mü (açık / fazla pay), beklenen şartlarla fark, zorunlu klozların
   anahtar kelimeleri slipte geçiyor mu.

Kullanım:
    python agent.py                                                         # örnek Cat XL slipi
    python agent.py --slip slip.pdf --beklenen beklenen.csv --klozlar zorunlu_klozlar.csv
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
RENK = {"Hata": "FDE2E1", "Dikkat": "FFF4CE", "Uygun": "E3F5E1", "Bilgi": "E8F0FE"}
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1 or (s.count(".") == 1 and len(s.split(".")[1]) == 3):
        s = s.replace(".", "")
    try:
        return Decimal(s) if s not in ("", "-", ".") else None
    except InvalidOperation:
        return None


def sayilar(metin: str) -> set[Decimal]:
    temiz = re.sub(r"\b\d{1,2}\.\d{1,2}\.\d{4}\b", " ", str(metin))
    return {v for x in re.findall(r"\d+(?:[.,]\d+)*", temiz) if (v := para(x)) is not None}


def tarihler(metin: str) -> set[str]:
    return set(re.findall(r"\b\d{1,2}\.\d{1,2}\.\d{4}\b", str(metin)))


def tablo(yol: Path) -> list[list[str]]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, read_only=True, data_only=True)
        s = [[str(c or "").strip() for c in r] for r in wb.active.iter_rows(values_only=True)]
        wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=metin.split("\n")[0].count)))
    return [r for r in s if any(c.strip() for c in r)]


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

ALANLAR = {"tur": "Tür", "tur_detay": "Tür ayrıntısı", "cedant": "Reasürans ettiren", "baslangic": "Başlangıç", "bitis": "Bitiş", "kapsam": "Kapsam",
           "limit": "Limit", "oncelik": "Öncelik", "para_birimi": "Para Birimi", "reinstatement_adedi": "Reinstatement Adedi",
           "reinstatement_kosulu": "Reinstatement koşulu", "mdp": "MDP", "prim_orani": "Prim Oranı %", "saat_klozu": "Saat Klozu",
           "bildirim_suresi": "Hasar bildirimi", "uyusmazlik": "Uyuşmazlık"}
SAYISAL = {"limit", "oncelik", "mdp", "prim_orani", "reinstatement_adedi"}
SEMA = {"type": "object", "properties": {
    **{a: S for a in ALANLAR if a != "tur"}, "tur": {"type": "string", "enum": ["Kotapar", "Eksedan", "Risk XL", "Cat XL", "Stop Loss", "Fakültatif", "Diğer"]},
    "istisnalar": {"type": "array", "items": S},
    "reasurorler": {"type": "array", "items": {"type": "object", "properties": {"ad": S, "pay": S, "lider": {"type": "boolean"}},
                                               "required": ["ad", "pay", "lider"], "additionalProperties": False}},
    "notlar": {"type": "array", "items": S}},
    "required": [*ALANLAR, "istisnalar", "reasurorler", "notlar"], "additionalProperties": False}


def denetle(c: dict, metin: str) -> list[str]:
    sorun, say, tar = [], sayilar(metin), tarihler(metin)
    for a in ("baslangic", "bitis"):
        if c[a].strip() and c[a].strip() not in tar:
            sorun.append(f"{ALANLAR[a]} '{c[a]}' slipte bu biçimde yazmıyor (GG.AA.YYYY)")
    for a in SAYISAL:
        v = c[a].strip()
        if v and (para(v) is None or para(v) not in say):
            sorun.append(f"{ALANLAR[a]} '{v}' slipte yazmıyor (hesaplamayın; slipte yoksa boş bırakın)")
    for a in ("cedant", "kapsam", "saat_klozu", "uyusmazlik"):
        if c[a].strip() and not (sayilar(c[a]) <= say):
            sorun.append(f"{ALANLAR[a]}: slipte olmayan sayı var ('{c[a][:50]}')")
    if c["cedant"].strip() and not metinde_var(c["cedant"], metin):
        sorun.append(f"Reasürans ettiren '{c['cedant']}' slipte aynen geçmiyor")
    for i, r in enumerate(c["reasurorler"], 1):
        if not metinde_var(r["ad"], metin):
            sorun.append(f"Reasürör {i}: '{r['ad']}' slipte aynen geçmiyor")
        if para(r["pay"]) is None or para(r["pay"]) not in say:
            sorun.append(f"Reasürör {i}: pay '{r['pay']}' slipte yazmıyor")
    if not c["reasurorler"] and not c["notlar"]:
        sorun.append("Reasürör listesi boş; slipte reasürör / pay yoksa bunu notlara yazın")
    return sorun


def karsilastir(c: dict, beklenen: list[tuple[str, str]]) -> list[dict]:
    """Kod: beklenen şartlarla slipteki değer. Sayısal beklenen değer slipteki değerin sayıları arasında olmalı."""
    out = []
    for kalem, bek in beklenen:
        k = katla(kalem)
        alan = next((a for a, ad in ALANLAR.items() if katla(ad) == k or katla(ad).startswith(k) or k.startswith(katla(ad))), None)
        if alan is None:
            out.append({"kalem": kalem, "slip": "", "beklenen": bek, "durum": "Bilgi", "not": "Bu kalem slip özetinde yok; elle kontrol edin"})
            continue
        v = c.get(alan, "")
        tarih = re.fullmatch(r"\d{1,2}\.\d{1,2}\.\d{4}", bek.strip())
        b_say = para(bek) if re.fullmatch(r"[\d.,\s%]+", bek.strip()) and not tarih else None
        if not v.strip():
            durum, nt = "Dikkat", "Slipte bulunamadı"
        elif tarih:
            durum, nt = ("Uygun", "") if bek.strip() in tarihler(v) else ("Hata", f"Beklenen {bek}, slipte {v}")
        elif b_say is not None:
            durum, nt = ("Uygun", "") if b_say in sayilar(v) else ("Hata", f"Beklenen {bek}, slipte {v}")
        else:
            durum, nt = ("Uygun", "") if katla(bek) == katla(v) or katla(bek) in katla(v) else ("Hata", f"Beklenen '{bek}', slipte '{v}'")
        out.append({"kalem": kalem, "slip": v, "beklenen": bek, "durum": durum, "not": nt})
    return out


def kontroller(c: dict, metin: str, klozlar: list[tuple[str, list[str]]]) -> list[tuple[str, str]]:
    b = []
    paylar = [para(r["pay"]) for r in c["reasurorler"]]
    if c["reasurorler"] and None not in paylar:
        top = sum(paylar, Decimal(0))
        yuzde = lambda x: f"%{x.normalize():f}".replace(".", ",")  # noqa: E731
        if top < 100:
            b.append(("Hata", f"Reasürör payları toplamı {yuzde(top)}: {yuzde(100 - top)} yerleşmemiş (açık pay)"))
        elif top > 100:
            b.append(("Hata", f"Reasürör payları toplamı {yuzde(top)}: {yuzde(top - 100)} fazla yerleşmiş (signing down gerekir)"))
    lider = [r["ad"] for r in c["reasurorler"] if r["lider"]]
    if c["reasurorler"] and len(lider) != 1:
        b.append(("Dikkat", f"Lider reasürör {'belirtilmemiş' if not lider else 'birden fazla: ' + ', '.join(lider)}"))
    km = katla(metin)
    for ad, kelimeler in klozlar:
        if not any(katla(k) and katla(k) in km for k in kelimeler):
            b.append(("Dikkat", f"Zorunlu kloz bulunamadı: {ad} (aranan: {', '.join(kelimeler)})"))
    lim, onc = para(c["limit"]), para(c["oncelik"])
    if c["tur"] in ("Cat XL", "Risk XL") and (lim is None or onc is None):
        b.append(("Dikkat", "Hasar fazlası anlaşmasında limit ya da öncelik bulunamadı"))
    if c["tur"] == "Cat XL" and not c["saat_klozu"].strip():
        b.append(("Dikkat", "Cat XL anlaşmasında saat klozu bulunamadı"))
    return b


def uret(metin: str, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    mesaj = f"<slip>\n{metin}\n</slip>"
    if duzeltme:
        mesaj += ("\n<duzeltme>\nÖnceki çıkarım: " + json.dumps(onceki, ensure_ascii=False) + "\nKodun bulduğu sorunlar:\n"
                  + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")
    return llm.json_iste(sistem, mesaj, SEMA)


def duzeltme_dongusu(uret_f, denetle_f, tur: int) -> tuple[dict, list[str], int]:
    taslak = uret_f(None, None)
    kalan = denetle_f(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret_f(kalan, taslak)
        kalan = denetle_f(taslak)
    return taslak, kalan, yapilan


# ----------------------------------------------------------------------------
# Çıktı
# ----------------------------------------------------------------------------

def _baslik(ws, basliklar, genislik) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font, h.alignment = BASLIK_DOLGU, BASLIK_YAZI, UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def rapor_yaz(cikti: Path, slip: str, c: dict, kars: list[dict], kont: list[tuple[str, str]], kalan: list[str], tur: int) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    _baslik(ws, ["Şart", "Slipteki Değer"], (24, 90))
    for a, ad in ALANLAR.items():
        ws.append([ad, c[a]])
        ws.cell(ws.max_row, 2).fill = AI_DOLGU
    ws.append(["İstisnalar", "\n".join(f"• {x}" for x in c["istisnalar"])])
    ws.append(["Model notları", "\n".join(f"• {x}" for x in c["notlar"])])
    for r in ws.iter_rows(min_row=2):
        for h in r:
            h.alignment = UST
    r = wb.create_sheet("Reasürörler")
    _baslik(r, ["Reasürör", "Pay %", "Lider"], (34, 10, 8))
    for x in c["reasurorler"]:
        r.append([x["ad"], None if para(x["pay"]) is None else float(para(x["pay"])), "Evet" if x["lider"] else ""])
    r.append(["Toplam", float(sum((para(x["pay"]) or Decimal(0) for x in c["reasurorler"]), Decimal(0))), ""])
    r.cell(r.max_row, 1).font = Font(bold=True)
    k = wb.create_sheet("Karşılaştırma")
    _baslik(k, ["Kalem", "Slipteki Değer", "Beklenen", "Durum", "Not"], (24, 40, 20, 9, 50))
    for x in kars:
        k.append([x["kalem"], x["slip"], x["beklenen"], x["durum"], x["not"]])
        k.cell(k.max_row, 4).fill = PatternFill("solid", fgColor=RENK[x["durum"]])
    b = wb.create_sheet("Bulgular")
    _baslik(b, ["Önem", "Bulgu"], (9, 100))
    for o, m in kont + [(x["durum"], f"{x['kalem']}: {x['not']}") for x in kars if x["durum"] in ("Hata", "Dikkat")]:
        b.append([o, m])
        b.cell(b.max_row, 1).fill = PatternFill("solid", fgColor=RENK[o])
    kt = wb.create_sheet("Kontrol")
    kt.append([f"Slip: {slip} · Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    kt.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    kt.cell(2, 1).fill = KIRMIZI if kalan else YESIL
    kt.append(["Önemli: Özet yapay zekâ ile çıkarılmıştır (mor); bağlayıcı olan slip / anlaşma metninin kendisidir. Pay toplamı, beklenen "
               "şartlarla fark ve zorunlu kloz kontrolü kodla yapılmıştır; kloz kontrolü anahtar kelime aramasıdır, klozun içeriğini "
               "değerlendirmez. İmzalamadan önce reasürans müdürü ve hukuk birimiyle kontrol edin."])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(slip: Path, cikti: Path, beklenen_yolu: Path | None = None, kloz_yolu: Path | None = None, tur: int = 2, evet: bool = False) -> dict:
    if not slip.exists():
        raise belge.BelgeHatasi(f"{slip}: slip bulunamadı.")
    metin = llm.maskele(belge.metin_oku(slip))
    beklenen = [(r[0].strip(), r[1].strip()) for r in tablo(beklenen_yolu)[1:] if len(r) > 1 and r[0].strip()] if beklenen_yolu else []
    klozlar = [(r[0].strip(), [k.strip() for k in r[1].split(",") if k.strip()]) for r in tablo(kloz_yolu)[1:] if len(r) > 1] if kloz_yolu else []
    print(f"[OK] Slip {len(metin):,} karakter · beklenen şart: {len(beklenen)} · zorunlu kloz: {len(klozlar)}")
    llm.onay_al(f"'{slip.name}' slip metni gönderilecek (telefon, e-posta, IBAN maskeli). Beklenen şartlar ve kloz listesi gönderilmez.", evet)
    c, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(metin, d, o), lambda x: denetle(x, metin), tur)
    kars = karsilastir(c, beklenen)
    kont = kontroller(c, metin, klozlar)
    rapor_yaz(cikti, slip.name, c, kars, kont, kalan, yapilan)
    return {"c": c, "karsilastirma": kars, "kontroller": kont, "kalan": kalan, "tur": yapilan}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Reasürans slipinden ana şartları çıkarır, beklenen şartlarla ve zorunlu klozlarla karşılaştırır.")
    p.add_argument("--slip", type=Path, default=ORNEK / "cat_xl_slip_2027.txt", help="Slip / anlaşma metni (.pdf, .docx, .txt)")
    p.add_argument("--beklenen", type=Path, default=None, help="Beklenen şartlar (.csv/.xlsx: Kalem; Beklenen)")
    p.add_argument("--klozlar", type=Path, default=None, help="Zorunlu klozlar (.csv/.xlsx: Kloz; Anahtar Kelimeler (virgülle))")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "slip_ozeti.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    ornek = a.slip == ORNEK / "cat_xl_slip_2027.txt"
    try:
        for y in (a.beklenen, a.klozlar):
            if y is not None and not y.exists():
                raise llm.LLMHatasi(f"{y}: bulunamadı.")
        s = calistir(a.slip, a.cikti, a.beklenen or (ORNEK / "beklenen.csv" if ornek else None), a.klozlar or (ORNEK / "zorunlu_klozlar.csv" if ornek else None),
                     a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    c = s["c"]
    print(f"[OK] {c['tur']} · {c['baslangic']} – {c['bitis']} · limit {c['limit']} xs {c['oncelik']} · {len(c['reasurorler'])} reasürör")
    for o, m in s["kontroller"] + [(x["durum"], f"{x['kalem']}: {x['not']}") for x in s["karsilastirma"] if x["durum"] == "Hata"]:
        print(f"[!] {o}: {m}")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
