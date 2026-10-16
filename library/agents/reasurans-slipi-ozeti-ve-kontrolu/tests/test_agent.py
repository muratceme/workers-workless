"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
SLIP = ORNEK / "cat_xl_slip_2027.txt"
METIN = SLIP.read_text(encoding="utf-8")
CEVAP = json.loads((Path(__file__).parent / "sahte_cevap.json").read_text(encoding="utf-8"))
MESAJLAR = []


def hatali() -> dict:
    """Sık model hataları: limit + önceliği toplamak, reasürör adını tamamlamak, tarihi farklı biçimde yazmak."""
    c = copy.deepcopy(CEVAP)
    c["limit"] = "450.000.000"
    c["reasurorler"][0]["ad"] = "Kurgu Reinsurance AG"
    c["bitis"] = "31/12/2027"
    return c


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return copy.deepcopy(CEVAP if "<duzeltme>" in kullanici else hatali())


def beklenen():
    return [(r[0], r[1]) for r in agent.tablo(ORNEK / "beklenen.csv")[1:]]


def klozlar():
    return [(r[0], [k.strip() for k in r[1].split(",")]) for r in agent.tablo(ORNEK / "zorunlu_klozlar.csv")[1:]]


class KodTesti(unittest.TestCase):
    def test_denetim(self):
        self.assertEqual(agent.denetle(CEVAP, METIN), [])
        s = agent.denetle(hatali(), METIN)
        self.assertEqual(len(s), 3)
        for parca in ("Limit '450.000.000' slipte yazmıyor", "'Kurgu Reinsurance AG' slipte aynen geçmiyor", "Bitiş '31/12/2027'"):
            self.assertTrue(any(parca in x for x in s), parca)

    def test_karsilastirma(self):
        k = {x["kalem"]: (x["durum"], x["not"]) for x in agent.karsilastir(CEVAP, beklenen())}
        self.assertEqual(k["MDP"], ("Hata", "Beklenen 18.000.000, slipte 16.500.000"))
        self.assertEqual({a: d for a, (d, _) in k.items() if d != "Uygun"}, {"MDP": "Hata"})
        c = copy.deepcopy(CEVAP) | {"saat_klozu": "", "tur": "Risk XL"}
        k = {x["kalem"]: x["durum"] for x in agent.karsilastir(c, beklenen() + [("Komisyon", "%25")])}
        self.assertEqual((k["Tür"], k["Saat Klozu (gün / saat)"], k["Komisyon"]), ("Hata", "Dikkat", "Bilgi"))

    def test_kontroller(self):
        b = agent.kontroller(CEVAP, METIN, klozlar())
        self.assertEqual(b, [("Hata", "Reasürör payları toplamı %95: %5 yerleşmemiş (açık pay)"),
                             ("Dikkat", "Zorunlu kloz bulunamadı: Siber istisnası / LMA siber klozu (aranan: siber, cyber)"),
                             ("Dikkat", "Zorunlu kloz bulunamadı: Yaptırım klozu (aranan: yaptırım, sanction)")])
        c = copy.deepcopy(CEVAP)
        c["reasurorler"].append({"ad": "Deneme Re SE", "pay": "7,5", "lider": True})
        b = agent.kontroller(c, METIN, [])
        self.assertEqual(b, [("Hata", "Reasürör payları toplamı %102,5: %2,5 fazla yerleşmiş (signing down gerekir)"),
                             ("Dikkat", "Lider reasürör birden fazla: Kurgu Re AG, Deneme Re SE")])
        self.assertIn("saat klozu", agent.kontroller(CEVAP | {"saat_klozu": ""}, METIN, [])[-1][1])


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ornek_ve_duzeltme(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "s.xlsx"
            self.assertEqual(agent.main(["--evet", "--cikti", str(cikti)]), 0)
            self.assertEqual(len(MESAJLAR), 2)
            self.assertIn("450.000.000", MESAJLAR[1].split("<duzeltme>")[1])
            self.assertNotIn("18.000.000", MESAJLAR[0])                    # beklenen şartlar sağlayıcıya gitmez
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Özet", "Reasürörler", "Karşılaştırma", "Bulgular", "Kontrol"])
            r = [list(x) for x in wb["Reasürörler"].iter_rows(min_row=2, values_only=True)]
            self.assertEqual(r[-1], ["Toplam", 95, None])
            self.assertEqual(wb["Bulgular"].max_row, 5)
            self.assertIn("sorunsuz", wb["Kontrol"].cell(2, 1).value)

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(agent.main(["--slip", str(Path(tmp) / "yok.pdf"), "--evet"]), 1)
            self.assertEqual(agent.main(["--beklenen", str(Path(tmp) / "yok.csv"), "--evet"]), 1)
            s = agent.calistir(SLIP, Path(tmp) / "s.xlsx", None, None, tur=0, evet=True)
            self.assertEqual((len(s["kalan"]), s["karsilastirma"]), (3, []))


if __name__ == "__main__":
    unittest.main()
