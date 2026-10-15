"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from PIL import Image  # noqa: E402

V = KLASOR / "ornek_veri"
GOR, REF = V / "gorseller", V / "referans_renkler.csv"
MESAJLAR = []


def r(renk, rol, tanim, kullanim):
    return {"renk": renk, "rol": rol, "tanim": tanim, "kullanim": kullanim}


def cevap(dogru=True) -> dict:
    t = {"palet_adi": "Kıyı Kumu", "hikaye": "Kum, kil ve deniz tonlarından oluşan sakin bir yaz paleti.",
         "renkler": [r("R1", "Ana", "Sıcak, açık kum beji (ÖRN-102).", "Keten gömlek ve pantolonlarda zemin; paletin %36'sı."),
                     r("R2", "Ana", "Açık kırık beyaz.", "Tişört ve astarlarda."),
                     r("R3", "Destek", "Soğuk, koyu gece mavisi.", "Alt giyimde kontrast."),
                     r("R4", "Destek", "Yumuşak adaçayı yeşili.", "Örme ürünlerde."),
                     r("R5", "Vurgu", "Sıcak kil tonu.", "Baskı detayı ve aksesuarda.")],
         "kombinasyonlar": ["R1 + R3: zemin ve kontrast biye", "R2 + R5: yaz tişörtü baskısı"]}
    if not dogru:
        t["renkler"][0]["rol"] = "Vurgu"
        t["renkler"][4]["rol"] = "Ana"
        t["renkler"][1]["rol"] = "Destek"
        t["renkler"][2]["tanim"] = "Sıcak gece mavisi (19-4052)."
        t["renkler"][1]["kullanim"] = "Paletin %40'ı."
        del t["renkler"][3]
    return t


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return cevap("<duzeltme>" in kullanici)


def renkler():
    x = agent.renkleri_cikar(agent.gorselleri_bul(GOR), 6, 8)
    ref = agent.referans_oku(REF)
    agent.eslestir(x, ref, 3)
    return x, ref


class KodTesti(unittest.TestCase):
    def test_renk_bilimi(self):
        self.assertAlmostEqual(agent.de2000((50, 2.6772, -79.7751), (50, 0, -82.7485)), 2.0425, places=4)  # Sharma vd. test verisi
        self.assertAlmostEqual(agent.de2000((50, -1.0, 2.0), (50, -1.0, 2.0)), 0.0)
        self.assertEqual((agent.rgb_hex(agent.hex_rgb("#c1784f")), agent.nitelik((46, 74, 98))["sicaklik"], agent.nitelik((128, 128, 128))["sicaklik"]),
                         ("#C1784F", "soğuk", "nötr"))

    def test_cikarim(self):
        x, _ = renkler()
        self.assertEqual([(y["ref"]["kod"], y["oran"]) for y in x],
                         [("ÖRN-102", 36), ("ÖRN-101", 22), ("ÖRN-105", 15), ("ÖRN-107", 15), ("ÖRN-103", 12)])
        self.assertTrue(all(y["yakin"] for y in x))
        self.assertEqual((x[0]["sicaklik"], x[2]["aciklik"]), ("sıcak", "koyu"))

    def test_denetim(self):
        x, ref = renkler()
        self.assertEqual(agent.denetle(cevap(), x, ref), [])
        s = agent.denetle(cevap(False), x, ref)
        for parca in ("R3: 'sıcak'", "R4: renge rol", "ana renk (R5)", "R1: vurgu", "'19-4052'", "yüzde: 40"):
            self.assertTrue(any(parca in y for y in s), (parca, s))
        self.assertEqual(len(s), 6, s)


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "p.xlsx"
            s = agent.calistir(GOR, REF, cikti, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            self.assertIn("R1 | #D4C2A7 | pay %36 | sıcak, açık | en yakın referans: ÖRN-102 Kum Beji", MESAJLAR[0])
            ws = load_workbook(cikti)["Palet"]
            self.assertEqual((ws["C2"].value, ws["B2"].fill.fgColor.rgb[-6:], ws["K2"].value), ("#D4C2A7", "D4C2A7", "Ana"))
            self.assertIn("Kıyı Kumu", s["kart"].read_text(encoding="utf-8"))

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(GOR, REF, Path(tmp) / "p.xlsx", tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 6)
            tek = Path(tmp) / "tek.png"
            Image.new("RGB", (40, 40), (31, 42, 68)).save(tek)
            self.assertEqual(agent.main(["--evet", "--gorseller", str(tek), "--renk", "3", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            bozuk = Path(tmp) / "bozuk.jpg"
            bozuk.write_bytes(b"jpeg degil")
            self.assertEqual(agent.main(["--evet", "--gorseller", str(bozuk)]), 1)
            self.assertEqual(agent.main(["--evet", "--renk", "20"]), 1)


if __name__ == "__main__":
    unittest.main()
