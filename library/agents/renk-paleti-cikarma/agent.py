"""
Renk Paleti Çıkarma — Gerisi Bende AI Agent
Tasarım ve Koleksiyon › Moda Tasarımcısı

1. Kod görsellerdeki baskın renkleri çıkarır (Pillow ile renk azaltma; yakın renkler CIEDE2000 ΔE < --birlestir ile
   birleştirilir) ve her rengin görsellerdeki payını hesaplar. Görseller modele gönderilmez; yalnız renk değerleri gönderilir.
2. Kod her rengi kullanıcının verdiği renk referans tablosundaki (ör. lisanslı Pantone TCX listesi: kod; ad; HEX) en
   yakın renkle eşleştirir (CIEDE2000) ve ΔE > --esik ise "yakın eşleşme yok" der. Renk sıcaklığı (sıcak / soğuk /
   nötr) ve açıklık (açık / orta / koyu) da kodla hesaplanır.
3. Model paletin adını ve hikâyesini, her rengin rolünü (Ana / Destek / Vurgu), kullanım önerisini ve kombinasyon
   notlarını yazar.
4. Kod denetler:
   - Her renk tam bir rol almış mı; en az bir "Ana" renk var mı ve ana renk en büyük paylı iki renkten biri mi; vurgu rengi
     ana renkten daha az paylı mı
   - Metindeki renk kodları eşleşen kodlardan mı (kod uydurulamaz); yüzdeler kodun hesabıyla aynı mı
   - "Sıcak", "soğuk", "açık", "koyu" nitelemeleri kodun hesabıyla çelişiyor mu
   Sorun varsa modele geri gönderilir (en fazla --tur kez).
5. Kod Excel (renkli hücreler) ve HTML palet kartı üretir.

Kullanım:
    python agent.py                                                        # örnek (kurgusal) ilham görselleri
    python agent.py --gorseller ilham/ --referans pantone_tcx.csv --renk 8
"""
from __future__ import annotations

import argparse
import colorsys
import csv
import html
import json
import math
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, UnidentifiedImageError

import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
UZANTI = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tif", ".tiff"}
REF = {"kod": ("Kod", "Renk Kodu", "TCX"), "ad": ("Ad", "Renk Adı", "İsim"), "hex": ("HEX", "Hex Kodu", "Renk")}
ROLLER = ["Ana", "Destek", "Vurgu"]
NITELIK = {"sicak": ("sicak", "sıcak"), "soguk": ("soguk", "soğuk"), "acik": ("acik", "açık"), "koyu": ("koyu", "koyu")}

SEMA = {
    "type": "object",
    "properties": {
        "palet_adi": S, "hikaye": S,
        "renkler": {"type": "array", "items": {"type": "object", "properties": {
            "renk": S, "rol": {"type": "string", "enum": ROLLER}, "tanim": S, "kullanim": S},
            "required": ["renk", "rol", "tanim", "kullanim"], "additionalProperties": False}},
        "kombinasyonlar": {"type": "array", "items": S},
    },
    "required": ["palet_adi", "hikaye", "renkler", "kombinasyonlar"],
    "additionalProperties": False,
}


# ---------------------------------------------------------------------------- renk bilimi
def hex_rgb(h: str) -> tuple[int, int, int]:
    h = str(h).strip().lstrip("#")
    if not re.fullmatch(r"[0-9a-fA-F]{6}", h):
        raise llm.LLMHatasi(f"HEX renk okunamadı: '{h}'")
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))


def rgb_hex(rgb) -> str:
    return "#%02X%02X%02X" % tuple(int(round(c)) for c in rgb)


def rgb_lab(rgb) -> tuple[float, float, float]:
    """sRGB (D65) → CIELAB."""
    def lin(c):
        c /= 255
        return c / 12.92 if c <= 0.04045 else ((c + 0.055) / 1.055) ** 2.4
    r, g, b = (lin(float(c)) for c in rgb)
    x = (0.4124 * r + 0.3576 * g + 0.1805 * b) / 0.95047
    y = 0.2126 * r + 0.7152 * g + 0.0722 * b
    z = (0.0193 * r + 0.1192 * g + 0.9505 * b) / 1.08883
    f = lambda t: t ** (1 / 3) if t > 0.008856 else 7.787 * t + 16 / 116
    fx, fy, fz = f(x), f(y), f(z)
    return 116 * fy - 16, 500 * (fx - fy), 200 * (fy - fz)


def de2000(l1, l2) -> float:
    """CIEDE2000 renk farkı (Sharma vd., 2005)."""
    L1, a1, b1 = l1
    L2, a2, b2 = l2
    c1, c2 = math.hypot(a1, b1), math.hypot(a2, b2)
    cm = (c1 + c2) / 2
    g = 0.5 * (1 - math.sqrt(cm ** 7 / (cm ** 7 + 25 ** 7)))
    a1p, a2p = a1 * (1 + g), a2 * (1 + g)
    c1p, c2p = math.hypot(a1p, b1), math.hypot(a2p, b2)
    h1p = math.degrees(math.atan2(b1, a1p)) % 360
    h2p = math.degrees(math.atan2(b2, a2p)) % 360
    dL, dC = L2 - L1, c2p - c1p
    dh = 0 if c1p * c2p == 0 else (h2p - h1p if abs(h2p - h1p) <= 180 else h2p - h1p - 360 if h2p > h1p else h2p - h1p + 360)
    dH = 2 * math.sqrt(c1p * c2p) * math.sin(math.radians(dh / 2))
    Lm, cmp_ = (L1 + L2) / 2, (c1p + c2p) / 2
    if c1p * c2p == 0:
        hm = h1p + h2p
    elif abs(h1p - h2p) <= 180:
        hm = (h1p + h2p) / 2
    else:
        hm = (h1p + h2p + 360) / 2 if h1p + h2p < 360 else (h1p + h2p - 360) / 2
    t = 1 - 0.17 * math.cos(math.radians(hm - 30)) + 0.24 * math.cos(math.radians(2 * hm)) + 0.32 * math.cos(math.radians(3 * hm + 6)) \
        - 0.20 * math.cos(math.radians(4 * hm - 63))
    sl = 1 + 0.015 * (Lm - 50) ** 2 / math.sqrt(20 + (Lm - 50) ** 2)
    sc, sh = 1 + 0.045 * cmp_, 1 + 0.015 * cmp_ * t
    rt = -2 * math.sqrt(cmp_ ** 7 / (cmp_ ** 7 + 25 ** 7)) * math.sin(math.radians(60 * math.exp(-(((hm - 275) / 25) ** 2))))
    return math.sqrt((dL / sl) ** 2 + (dC / sc) ** 2 + (dH / sh) ** 2 + rt * (dC / sc) * (dH / sh))


def nitelik(rgb) -> dict:
    L, a, b = rgb_lab(rgb)
    c = math.hypot(a, b)
    h = math.degrees(math.atan2(b, a)) % 360
    if c < 10:
        sicaklik = "nötr"
    elif h < 100 or h >= 330:
        sicaklik = "sıcak"
    elif 160 <= h < 300:
        sicaklik = "soğuk"
    else:
        sicaklik = "nötr"
    return {"sicaklik": sicaklik, "aciklik": "açık" if L >= 70 else ("koyu" if L < 35 else "orta"), "L": L, "C": c}


# ---------------------------------------------------------------------------- çıkarım
def gorselleri_bul(yol: Path) -> list[Path]:
    if yol.is_file():
        return [yol]
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: görsel ya da klasör bulunamadı.")
    g = sorted(p for p in yol.iterdir() if p.suffix.lower() in UZANTI)
    if not g:
        raise llm.LLMHatasi(f"{yol}: görsel yok ({', '.join(sorted(UZANTI))}).")
    return g


def renkleri_cikar(gorseller: list[Path], adet: int, birlestir: float) -> list[dict]:
    havuz = []  # (rgb, pay) — her görselin ağırlığı eşit
    for g in gorseller:
        try:
            with Image.open(g) as im:
                im = im.convert("RGB")
                im.thumbnail((200, 200))
                q = im.quantize(colors=adet * 2, method=Image.Quantize.MEDIANCUT)
                pal = q.getpalette()
                say = q.getcolors()
        except (UnidentifiedImageError, OSError) as h:
            raise llm.LLMHatasi(f"{g.name}: görsel açılamadı ({h}).")
        toplam = sum(n for n, _ in say)
        havuz += [((pal[i * 3], pal[i * 3 + 1], pal[i * 3 + 2]), n / toplam / len(gorseller)) for n, i in say]
    kumeler = []  # ağırlıklı ortalama ile birleştir
    for rgb, pay in sorted(havuz, key=lambda x: -x[1]):
        lab = rgb_lab(rgb)
        k = next((k for k in kumeler if de2000(k["lab"], lab) < birlestir), None)
        if k:
            t = k["pay"] + pay
            k["rgb"] = tuple((k["rgb"][j] * k["pay"] + rgb[j] * pay) / t for j in range(3))
            k["pay"], k["lab"] = t, rgb_lab(k["rgb"])
        else:
            kumeler.append({"rgb": rgb, "pay": pay, "lab": lab})
    kumeler = sorted(kumeler, key=lambda k: -k["pay"])[:adet]
    toplam = sum(k["pay"] for k in kumeler)
    out = []
    for i, k in enumerate(kumeler, 1):
        out.append({"id": f"R{i}", "hex": rgb_hex(k["rgb"]), "rgb": tuple(int(round(c)) for c in k["rgb"]), "lab": k["lab"],
                    "oran": round(k["pay"] / toplam * 100)} | nitelik(k["rgb"]))
    fark = 100 - sum(r["oran"] for r in out)
    if out:
        out[0]["oran"] += fark  # yuvarlama farkı en büyük paya
    return out


def referans_oku(yol: Path) -> list[dict]:
    out = []
    for r in kayitlar(yol, REF, ("kod", "hex")):
        if r["kod"]:
            rgb = hex_rgb(r["hex"])
            out.append({"kod": str(r["kod"]).strip(), "ad": str(r["ad"]).strip(), "hex": rgb_hex(rgb), "lab": rgb_lab(rgb)})
    if not out:
        raise llm.LLMHatasi(f"{yol.name}: referans tablosu boş.")
    return out


def eslestir(renkler: list[dict], ref: list[dict], esik: float) -> None:
    for r in renkler:
        en = min(ref, key=lambda x: de2000(r["lab"], x["lab"]))
        r["ref"], r["de"] = en, de2000(r["lab"], en["lab"])
        r["yakin"] = r["de"] <= esik


# ---------------------------------------------------------------------------- model ve denetim
def denetle(t: dict, renkler: list[dict], ref: list[dict]) -> list[str]:
    sorun = []
    r = {x["id"]: x for x in renkler}
    roller = {}
    for x in t["renkler"]:
        if x["renk"] not in r:
            sorun.append(f"{x['renk']}: böyle bir renk yok")
            continue
        if x["renk"] in roller:
            sorun.append(f"{x['renk']}: birden fazla rol almış")
        roller[x["renk"]] = x["rol"]
        tk = katla(x["tanim"] + " " + x["kullanim"])
        n = r[x["renk"]]
        for anahtar, (k, ad) in NITELIK.items():
            if re.search(rf"\b{k}\b", tk):
                if anahtar in ("sicak", "soguk") and n["sicaklik"] != ad:
                    sorun.append(f"{x['renk']}: '{ad}' denmiş ama renk kod hesabına göre {n['sicaklik']}")
                if anahtar in ("acik", "koyu") and n["aciklik"] != ad and not (anahtar == "acik" and n["aciklik"] == "orta" and n["L"] >= 60):
                    sorun.append(f"{x['renk']}: '{ad}' denmiş ama renk kod hesabına göre {n['aciklik']} tonda (L*={n['L']:.0f})")
    for k in r:
        if k not in roller:
            sorun.append(f"{k}: renge rol verilmemiş")
    anar = [k for k, v in roller.items() if v == "Ana"]
    ilk_iki = {x["id"] for x in sorted(renkler, key=lambda x: -x["oran"])[:2]}
    if not anar:
        sorun.append("Roller: 'Ana' renk yok")
    elif not set(anar) & ilk_iki:
        sorun.append(f"Roller: ana renk ({', '.join(anar)}) en büyük paylı iki renkten ({', '.join(sorted(ilk_iki))}) biri değil")
    for k, v in roller.items():
        if v == "Vurgu" and anar and r[k]["oran"] >= max(r[a]["oran"] for a in anar if a in r):
            sorun.append(f"{k}: vurgu rengi ana renkten daha büyük paylı (%{r[k]['oran']})")
    tum = " ".join([t["palet_adi"], t["hikaye"]] + t["kombinasyonlar"] + [x["tanim"] + " " + x["kullanim"] for x in t["renkler"]])
    eslesen = {x["ref"]["kod"] for x in renkler}
    for kod in {x["kod"] for x in ref} - eslesen:
        if re.search(rf"(?<![\w-]){re.escape(kod)}(?![\w-])", tum):
            sorun.append(f"Metin: {kod} paletteki renklerle eşleşmeyen bir referans kodu")
    for m in re.finditer(r"\b\d{2}-\d{4}\b", tum):
        if m.group(0) not in " ".join(eslesen):
            sorun.append(f"Metin: '{m.group(0)}' renk kodu referans eşleşmelerinde yok (kod uydurulamaz)")
    izinli = {str(x["oran"]) for x in renkler} | sayilar(" ".join(x["ref"]["kod"] + " " + x["hex"] for x in renkler))
    fazla = sorted({re.sub(r"\D", "", m) for m in re.findall(r"%\s*\d+|\d+\s*%", tum)} - izinli)
    if fazla:
        sorun.append(f"Metin: kodun hesabında olmayan yüzde: {', '.join(fazla)}")
    return list(dict.fromkeys(sorun))


def uret(renkler: list[dict], brif: str, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    satir = "\n".join(f"{x['id']} | {x['hex']} | pay %{x['oran']} | {x['sicaklik']}, {x['aciklik']} | en yakın referans: {x['ref']['kod']} {x['ref']['ad']} "
                      f"(ΔE {x['de']:.1f}{'' if x['yakin'] else ', yakın eşleşme yok'})" for x in renkler)
    mesaj = f"<brif>{brif or '(verilmedi)'}</brif>\n<renkler>\n{satir}\n</renkler>" + duzeltme_blogu(duzeltme, onceki)
    return llm.json_iste(sistem, mesaj, SEMA)


# ---------------------------------------------------------------------------- çıktı
def rapor_yaz(cikti: Path, t, renkler, gorseller, kalan, tur) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Palet"
    baslik(ws, ["Renk", "Örnek", "HEX", "RGB", "Pay %", "Sıcaklık", "Açıklık", "En Yakın Referans", "ΔE00", "Eşleşme", "Rol", "Tanım", "Kullanım"],
           (6, 10, 10, 14, 8, 10, 10, 30, 8, 14, 9, 40, 40))
    m = {x["renk"]: x for x in t["renkler"]}
    for x in renkler:
        y = m.get(x["id"], {})
        ws.append([x["id"], "", x["hex"], ", ".join(map(str, x["rgb"])), x["oran"], x["sicaklik"], x["aciklik"], f"{x['ref']['kod']} {x['ref']['ad']}",
                   round(x["de"], 1), "Yakın" if x["yakin"] else "Yakın eşleşme yok", y.get("rol", ""), y.get("tanim", ""), y.get("kullanim", "")])
        ws.cell(ws.max_row, 2).fill = PatternFill("solid", fgColor=x["hex"][1:])
        boya(ws, (10,), YESIL if x["yakin"] else SARI)
        boya(ws, (11, 12, 13), AI_DOLGU)
    son_not(ws, f"{t['palet_adi']}: {t['hikaye']}\nKombinasyonlar: " + " · ".join(t["kombinasyonlar"]) +
            f"\nKaynak görseller: {', '.join(g.name for g in gorseller)}\nÖnemli: Ekran renkleri cihaza göre değişir; renk onayını referans kartı / "
            "kumaş laboratuvar örneği (lab dip) ile ışık kabininde yapın.", kalan, tur)
    hizala(ws)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    kart = cikti.with_suffix(".html")
    kutular = "".join(
        f'<div class="k"><div class="s" style="background:{x["hex"]};flex:{max(x["oran"], 4)}"></div><b>{html.escape(m.get(x["id"], {}).get("rol", ""))}</b>'
        f'<span>{x["hex"]} · %{x["oran"]}</span><span>{html.escape(x["ref"]["kod"] + " " + x["ref"]["ad"])}</span></div>' for x in renkler)
    kart.write_text(f"""<!doctype html><html lang="tr"><head><meta charset="utf-8"><title>{html.escape(t['palet_adi'])}</title>
<style>body{{font-family:system-ui,sans-serif;margin:24px;background:#fff;color:#1d1d1f}}.p{{display:flex;gap:8px;flex-wrap:wrap}}
.k{{display:flex;flex-direction:column;width:150px;font-size:13px;gap:2px}}.s{{height:120px;border-radius:8px;border:1px solid #ddd}}</style></head>
<body><h1>{html.escape(t['palet_adi'])}</h1><p>{html.escape(t['hikaye'])}</p><div class="p">{kutular}</div>
<p><small>Renkler ekranda yaklaşık görünür; onay için fiziksel referans kartı kullanın.</small></p></body></html>""", encoding="utf-8")
    return kart


def calistir(gorsel_yol: Path, ref_yol: Path, cikti: Path, adet: int = 6, esik: float = 3.0, birlestir: float = 8.0, brif: str = "",
             tur: int = 2, evet: bool = False) -> dict:
    gorseller = gorselleri_bul(gorsel_yol)
    ref = referans_oku(ref_yol)
    renkler = renkleri_cikar(gorseller, adet, birlestir)
    eslestir(renkler, ref, esik)
    print(f"[OK] {len(gorseller)} görsel · {len(renkler)} baskın renk · referans tablosunda {len(ref)} renk")
    for x in renkler:
        print(f"    {x['id']} {x['hex']} %{x['oran']:<3} → {x['ref']['kod']} {x['ref']['ad']} (ΔE {x['de']:.1f})")
    llm.onay_al(f"{len(renkler)} rengin HEX değerleri, payları ve eşleşen referans adları gönderilecek; görseller gönderilmez.", evet)
    t, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(renkler, brif, d, o), lambda x: denetle(x, renkler, ref), tur)
    kart = rapor_yaz(cikti, t, renkler, gorseller, kalan, yapilan)
    return {"renkler": renkler, "taslak": t, "kalan": kalan, "tur": yapilan, "kart": kart}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    p = argparse.ArgumentParser(description="Görsellerden baskın renkleri çıkarıp referans tablosundaki (ör. Pantone TCX) en yakın kodlarla renk paleti oluşturur.")
    v = BURASI / "ornek_veri"
    p.add_argument("--gorseller", type=Path, default=v / "gorseller", help="Görsel dosyası ya da klasörü (jpg, png, webp …)")
    p.add_argument("--referans", type=Path, default=v / "referans_renkler.csv", help="Renk referans tablosu (CSV/Excel): Kod, Ad, HEX — ör. lisanslı Pantone TCX listeniz")
    p.add_argument("--renk", type=int, default=6, help="Paletteki renk sayısı (varsayılan 6)")
    p.add_argument("--esik", type=float, default=3.0, help="Yakın eşleşme için en fazla ΔE00 (varsayılan 3)")
    p.add_argument("--birlestir", type=float, default=8.0, help="Bu ΔE00'ın altındaki renkleri tek renk say (varsayılan 8)")
    p.add_argument("--brif", default="", help="Sezon / koleksiyon brifi (ör. 'İlkbahar-Yaz kadın keten koleksiyonu')")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "renk_paleti.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    if not 2 <= a.renk <= 12:
        print("[X] --renk 2 ile 12 arasında olmalı.")
        return 1
    try:
        s = calistir(a.gorseller, a.referans, a.cikti, a.renk, a.esik, a.birlestir, a.brif, a.tur, a.evet)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {s['taslak']['palet_adi']}")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Palet kartı: {s['kart'].resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
