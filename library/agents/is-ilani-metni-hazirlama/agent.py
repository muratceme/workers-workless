"""
İş İlanı Metni Hazırlama — Workers / Workless AI Agent
İnsan Kaynakları › İşe Alım Uzmanı

1. Kod, departmanın pozisyon talep formunu ayrımcı/riskli ifadeler için tarar (yaş, cinsiyet, medeni hâl,
   görünüş, sağlık, din, köken, sendika; dikkat: askerlik, ehliyet, uyruk, gereksiz kişisel veri).
2. Model, bulgular ışığında yayına hazır ilan yazar: ayrımcı şartları çıkarır ve listeler, işin gereği olan
   şartları gerekçesiyle yazar, kişilik özelliklerini yetkinliğe çevirir, KVKK yer tutucusu ekler.
3. Kod, üretilen ilanı yeniden tarar. "Yüksek" riskli ifade kalmışsa ya da "dikkat" ifadesi gerekçesiz
   kullanılmışsa taslak hata listesiyle modele geri gönderilir (en fazla --tur kez). Kalanlar raporda işaretlenir.

Dayanak: 4857 s. İş Kanunu md. 5; 6701 s. TİHEK Kanunu md. 3, 6 ve 7; 6356 s. Kanun md. 25; KVKK md. 4.
Tarama kural tabanlıdır ve hukuki değerlendirme yerine geçmez.

Kullanım:
    python agent.py                                              # örnek talep formuyla dener
    python agent.py --girdi talep_formu.docx
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import ayrimcilik
import belge
import llm

BURASI = Path(__file__).resolve().parent

SEMA = {
    "type": "object",
    "properties": {
        "baslik": {"type": "string"},
        "giris": {"type": "string"},
        "sorumluluklar": {"type": "array", "items": {"type": "string"}},
        "aranan_nitelikler": {"type": "array", "items": {"type": "string"}},
        "tercih_sebepleri": {"type": "array", "items": {"type": "string"}},
        "sunulanlar": {"type": "array", "items": {"type": "string"}},
        "calisma_bilgisi": {"type": "string"},
        "basvuru": {"type": "string"},
        "kisa_versiyon": {"type": "string"},
        "cikarilan_sartlar": {"type": "array", "items": {"type": "object", "properties": {"ifade": {"type": "string"}, "neden": {"type": "string"}},
                                                          "required": ["ifade", "neden"], "additionalProperties": False}},
        "dikkat_notlari": {"type": "array", "items": {"type": "object", "properties": {"ifade": {"type": "string"}, "gerekce": {"type": "string"}},
                                                       "required": ["ifade", "gerekce"], "additionalProperties": False}},
    },
    "required": ["baslik", "giris", "sorumluluklar", "aranan_nitelikler", "tercih_sebepleri", "sunulanlar", "calisma_bilgisi", "basvuru",
                 "kisa_versiyon", "cikarilan_sartlar", "dikkat_notlari"],
    "additionalProperties": False,
}


def ilan_metni(t: dict) -> str:
    bolumler = [t["baslik"], "", t["giris"], "", "Sorumluluklar", *[f"• {x}" for x in t["sorumluluklar"]], "", "Aranan nitelikler",
                *[f"• {x}" for x in t["aranan_nitelikler"]]]
    if t["tercih_sebepleri"]:
        bolumler += ["", "Tercih sebepleri", *[f"• {x}" for x in t["tercih_sebepleri"]]]
    if t["sunulanlar"]:
        bolumler += ["", "Sunduklarımız", *[f"• {x}" for x in t["sunulanlar"]]]
    bolumler += ["", t["calisma_bilgisi"], "", t["basvuru"]]
    return "\n".join(bolumler)


def denetle(t: dict) -> list[dict]:
    """Üretilen ilandaki riskli ifadeler: yüksek her zaman; dikkat yalnız gerekçesi yazılmamışsa."""
    metin = ilan_metni(t) + "\n" + t["kisa_versiyon"]
    gerekceli = " ".join(ayrimcilik.kucuk(d["ifade"] + " " + d["gerekce"]) for d in t["dikkat_notlari"])
    sorunlar = []
    for b in ayrimcilik.tara(metin):
        if b["seviye"] == "yüksek" or b["kategori"] == "Kişisel veri":       # gerekçeyle de kullanılmaz
            sorunlar.append(b)
        elif not any(k in gerekceli for k in (b["ifade"], b["kategori"].lower())):
            sorunlar.append({**b, "aciklama": b["aciklama"] + " — dikkat_notlari'nda gerekçe yok"})
    return sorunlar


def uret(talep: str, bulgular: list[dict], duzeltme: list[dict] | None = None, onceki: dict | None = None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    parcalar = [f"<talep_formu>\n{llm.maskele(talep)}\n</talep_formu>",
                "<tarayici_bulgulari>\n" + ("\n".join(f"- [{b['seviye']}] {b['kategori']}: '{b['ifade']}' ({b['aciklama']})" for b in bulgular)
                                            or "- Bulgu yok") + "\n</tarayici_bulgulari>"]
    if duzeltme:
        parcalar.append("<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False) + "\nKodun bulduğu sorunlar:\n"
                        + "\n".join(f"- {b['kategori']}: '{b['ifade']}' ({b['aciklama']})" for b in duzeltme) + "\n</duzeltme>")
    return llm.json_iste(sistem, "\n".join(parcalar), SEMA)


BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)


def rapor_yaz(cikti: Path, kaynak: Path, girdi_bulgulari: list[dict], t: dict, kalan: list[dict], tur: int) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "İlan"
    ws.append(["İlan metni (yayına hazır taslak)"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([ilan_metni(t)])
    ws.append([])
    ws.append(["Kısa versiyon (sosyal medya)"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.append([t["kisa_versiyon"]])
    ws.append([])
    ws.append(["Son kontrol: " + (ayrimcilik.ozet(kalan) if kalan else "Ayrımcı/riskli ifade bulunmadı")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append(["Önemli: Yapay zekâ taslağıdır; yayımlamadan önce İK ve gerekiyorsa hukuk birimi onayı alınmalıdır. "
               "Tarama kural tabanlıdır, hukuki değerlendirme yerine geçmez."])
    ws.column_dimensions["A"].width = 120
    for satir in ws.iter_rows():
        for h in satir:
            h.alignment = UST

    d = wb.create_sheet("Uygunluk")
    d.append(["Bölüm", "Seviye / Kategori", "İfade", "Açıklama / Gerekçe", "Onay"])
    for h in d[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for b in girdi_bulgulari:
        d.append(["Talep formunda bulunan", f"{b['seviye']} · {b['kategori']}", b["ifade"], b["aciklama"], ""])
        d.cell(d.max_row, 2).fill = KIRMIZI if b["seviye"] == "yüksek" else SARI
    for c in t["cikarilan_sartlar"]:
        d.append(["İlandan çıkarıldı", "", c["ifade"], c["neden"], ""])
    for c in t["dikkat_notlari"]:
        d.append(["Gerekçeyle tutuldu", "dikkat", c["ifade"], c["gerekce"], ""])
        d.cell(d.max_row, 2).fill = SARI
    for b in kalan:
        d.append(["İlanda KALAN sorun", f"{b['seviye']} · {b['kategori']}", b["ifade"], b["aciklama"], ""])
        d.cell(d.max_row, 1).fill = KIRMIZI
    for j, w in enumerate((24, 26, 34, 80, 10), 1):
        d.column_dimensions["ABCDE"[j - 1]].width = w
    for r in range(2, d.max_row + 1):
        d.cell(r, 5).fill = SARI
        for c in range(1, 6):
            d.cell(r, c).alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    md = cikti.with_suffix(".md")
    md.write_text(ilan_metni(t).replace("\n• ", "\n- ") + "\n\n---\n_Kısa versiyon:_ " + t["kisa_versiyon"] + "\n", encoding="utf-8")
    return md


def calistir(girdi: Path, cikti: Path, tur: int = 2, evet: bool = False) -> dict:
    talep = belge.metin_oku(girdi)
    bulgular = ayrimcilik.tara(talep)
    print(f"[OK] Talep formu tarandı: {sum(b['seviye'] == 'yüksek' for b in bulgular)} yüksek, "
          f"{sum(b['seviye'] == 'dikkat' for b in bulgular)} dikkat bulgusu")
    llm.onay_al("Pozisyon talep formu (e-posta/telefon maskeli) ve tarama bulguları ilan yazımı için gönderilecek.", evet)
    taslak = uret(talep, bulgular)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} riskli ifade modele geri gönderiliyor")
        taslak = uret(talep, bulgular, kalan, taslak)
        kalan = denetle(taslak)
    md = rapor_yaz(cikti, girdi, bulgular, taslak, kalan, yapilan)
    return {"girdi_bulgulari": bulgular, "taslak": taslak, "kalan": kalan, "tur": yapilan, "md": md}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Pozisyon talep formundan ayrımcı ifade içermeyen iş ilanı hazırlar.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "pozisyon_talep_formu.txt", help="Talep formu (.txt, .md, .docx, .pdf)")
    p.add_argument("--tur", type=int, default=2, help="Riskli ifade kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "is_ilani.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.cikti, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] İlan hazır · çıkarılan şart: {len(s['taslak']['cikarilan_sartlar'])} · gerekçeyle tutulan: {len(s['taslak']['dikkat_notlari'])}")
    for b in s["kalan"]:
        print(f"[!] Kalan: {b['kategori']} '{b['ifade']}'")
    print(f"[OK] Rapor: {a.cikti.resolve()} (+ {s['md'].name})")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
