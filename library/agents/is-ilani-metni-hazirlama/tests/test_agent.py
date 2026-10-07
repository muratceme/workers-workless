"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import ayrimcilik  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri" / "pozisyon_talep_formu.txt"
MESAJLAR = []


def taslak(nitelikler, dikkat=()):
    return {"baslik": "Saha Satış Temsilcisi (Ege Bölgesi)", "giris": "Örnek Gıda Sanayi A.Ş. ...", "sorumluluklar": ["Market ziyaretleri"],
            "aranan_nitelikler": list(nitelikler), "tercih_sebepleri": [], "sunulanlar": ["Şirket aracı"], "calisma_bilgisi": "Tam zamanlı",
            "basvuru": "[KVKK aydınlatma metni bağlantısı]", "kisa_versiyon": "Saha satış temsilcisi arıyoruz.",
            "cikarilan_sartlar": [{"ifade": "25-35 yaş arası", "neden": "Yaş ayrımcılığı"}],
            "dikkat_notlari": [{"ifade": i, "gerekce": g} for i, g in dikkat]}


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    if "<duzeltme>" not in kullanici:      # ilk taslak bilerek hatalı: "hoş görünümlü" kalmış, ehliyet gerekçesiz
        return taslak(["En az 2 yıl FMCG saha satış deneyimi", "Hoş görünümlü", "B sınıfı ehliyet"])
    return taslak(["En az 2 yıl FMCG saha satış deneyimi", "Müşteri ilişkilerinde güçlü iletişim", "B sınıfı ehliyet"],
                  [("B sınıfı ehliyet", "Sahada şirket aracıyla çalışılacağı için")])


class TaramaTesti(unittest.TestCase):
    def test_talep_formu(self):
        b = ayrimcilik.tara(ORNEK.read_text(encoding="utf-8"))
        kat = {(x["seviye"], x["kategori"]) for x in b}
        self.assertTrue({("yüksek", "Yaş"), ("yüksek", "Cinsiyet"), ("yüksek", "Görünüş"), ("dikkat", "Askerlik"), ("dikkat", "Ehliyet")} <= kat)

    def test_yanlis_alarm_yok(self):
        for t in ("Sağlıklı beslenme ürünleri", "dini bayramlarda vardiya", "güzel bir çalışma ortamı", "Excel bilgisi", "dinamik ekip"):
            self.assertEqual(ayrimcilik.tara(t), [], t)

    def test_denetim(self):
        self.assertEqual([b["kategori"] for b in agent.denetle(taslak(["Fotoğraflı CV gönderin"], [("Fotoğraf", "kurum kimliği")]))],
                         ["Kişisel veri"])                                # kişisel veri gerekçeyle de kabul edilmez
        self.assertEqual(agent.denetle(taslak(["B sınıfı ehliyet"], [("B sınıfı ehliyet", "Saha")])), [])


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "i.xlsx"
            s = agent.calistir(ORNEK, cikti, evet=True)
            self.assertEqual(len(MESAJLAR), 2)
            self.assertIn("<tarayici_bulgulari>", MESAJLAR[0])
            self.assertIn("hoş görünümlü", MESAJLAR[1].split("<duzeltme>")[1])
            self.assertEqual((s["kalan"], s["tur"]), ([], 1))
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["İlan", "Uygunluk"])
            self.assertIn("Müşteri ilişkilerinde", wb["İlan"]["A2"].value)
            self.assertIn("Ayrımcı/riskli ifade bulunmadı", wb["İlan"]["A7"].value)
            self.assertTrue(s["md"].exists())

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(ORNEK, Path(tmp) / "i.xlsx", tur=0, evet=True)
        self.assertEqual({b["kategori"] for b in s["kalan"]}, {"Görünüş", "Ehliyet"})


if __name__ == "__main__":
    unittest.main()
