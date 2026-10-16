"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import json
import os
import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
CEVAP = json.loads((Path(__file__).parent / "sahte_cevaplar.json").read_text(encoding="utf-8"))
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    """İlk turda S2 cevabı ödeme sözü verir ('bordroya eklenecektir'); düzeltmede doğru cevap döner."""
    MESAJLAR.append(kullanici)
    if "1.750 TL az yattı" in kullanici:
        return copy.deepcopy(CEVAP["S1"])
    if "10 saat fazla mesai" in kullanici:
        return copy.deepcopy(CEVAP["S2"] if "<duzeltme>" in kullanici else CEVAP["S2_hatali"])
    return copy.deepcopy(CEVAP["S3"])


def hazir():
    q, _ = agent.hazirla(agent.sorulari_oku(ORNEK / "sorular.txt"), agent.bordro_oku(ORNEK / "bordro.csv"))
    return {x["no"]: x for x in q}


class KodTesti(unittest.TestCase):
    def test_okuma(self):
        b = agent.bordro_oku(ORNEK / "bordro.csv")
        self.assertEqual(sorted(b), ["1001", "1002", "1003"])
        self.assertEqual((b["1003"]["ad"], b["1003"]["aylar"][10]["gun"], b["1003"]["aylar"][10]["eksik_neden"], b["1002"]["aylar"][9]["fazla_mesai"]),
                         ("Zeynep Kurgu", D(22), "Ücretsiz izin (8 gün)", D("4650.00")))
        k = agent.bordro_oku(Path(__file__).parent / "kod_blogu_cikti.xlsx")      # Brüt-Net Maaş Hesaplama çıktısı
        self.assertEqual((sorted(k)[0], k["Ayşe Yılmaz"]["aylar"][10]["net"], len(k["Ayşe Yılmaz"]["aylar"])), ("Ali Veli", D("37736.3"), 12))
        s = agent.sorulari_oku(ORNEK / "sorular.txt")
        self.assertEqual([(x["kim"], x["ay_metin"]) for x in s], [("1001", "Ekim"), ("1002", "Ekim"), ("Zeynep Kurgu", "Ekim")])
        self.assertEqual([agent.ay_coz(x) for x in ("Ekim", "ekim 2026", "10", "2026-10", "13", "x")], [10, 10, 10, 10, None, None])
        self.assertEqual((agent.calisan_bul(b, "zeynep kurgu"), agent.calisan_bul(b, "Zeynep"), agent.calisan_bul(b, "1002")), ("1003", None, "1002"))

    def test_analiz(self):
        q = hazir()
        self.assertEqual(q["S1"]["analiz"]["isaret"], [("Neden", "Vergi dilimi %20 → %27 (kümülatif matrah 425.000,00)")])
        self.assertEqual([t for t, _ in q["S2"]["analiz"]["isaret"]], ["Neden", "İnsan kontrolü", "İnsan kontrolü"])
        self.assertEqual([m for _, m in q["S3"]["analiz"]["isaret"]],
                         ["Gün sayısı 30 → 22 (Ücretsiz izin (8 gün))", "Asgari ücret GV istisnası 5.615,10 → 4.117,74"])
        self.assertIn("Net ücret | 32.762,20 | 24.025,60 | -8.736,60", q["S3"]["hesap"])
        self.assertTrue({D("8736.60"), D("1750.00"), D(27), D(8)} <= q["S3"]["izinli"] | q["S1"]["izinli"])
        b = agent.bordro_oku(ORNEK / "bordro.csv")
        a = agent.analiz(b["1001"], 9, "Maaşımda hata var")
        self.assertEqual([t for t, _ in a["isaret"]], ["Bilgi"])                    # önceki ay yok; karşılaştırma yapılamaz
        a = agent.analiz(b["1001"] | {"aylar": {9: b["1001"]["aylar"][9], 10: b["1001"]["aylar"][9]}}, 10, "maaşım yanlış yattı")
        self.assertEqual(a["isaret"], [("İnsan kontrolü", "Çalışan hata bildiriyor ama iki ay arasında bordroda fark yok")])

    def test_maskeleme_ve_denetim(self):
        q = hazir()
        self.assertTrue(q["S1"]["maskeli"].endswith("Teşekkürler, [ÇALIŞAN]"))
        self.assertEqual([agent.denetle(q[n], CEVAP[n]) for n in ("S1", "S2", "S3")], [[], [], []])
        s = agent.denetle(q["S2"], CEVAP["S2_hatali"])
        self.assertEqual(s, ["S2: cevapta ödeme / düzeltme sözü var ('bordroya eklenecektir'); söz vermeyin, kontrol edilip bilgi verileceğini yazın"])
        c = CEVAP["S1"] | {"cevap": "Sayın [ÇALIŞAN], vergi 1.800,00 TL arttı; fark iade edilecektir.", "insan_kontrolu": True, "kontrol_nedeni": ""}
        s = agent.denetle(q["S1"], c)
        self.assertEqual(len(s), 5)
        for parca in ("olmayan sayı: 1800", "iade edilecektir", "neden yazılmamış", "hitapla başlamasın", "yer tutucusu"):
            self.assertTrue(any(parca in x for x in s), parca)
        s = agent.denetle(q["S2"], CEVAP["S2"] | {"insan_kontrolu": False})
        self.assertIn("insan_kontrolu true olmalı", s[0])
        self.assertFalse(agent.SOZ.search("Ödenecek gelir vergisi 2.884,90 TL'den 4.634,90 TL'ye çıktı; size bilgi verilecektir."))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ornek_ve_duzeltme(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "c.xlsx"
            s = agent.calistir(ORNEK / "bordro.csv", ORNEK / "sorular.txt", cikti, imza="Bordro Birimi", evet=True)
            self.assertEqual(len(MESAJLAR), 4)                         # 3 soru + yalnız S2 için düzeltme
            self.assertIn("bordroya eklenecektir", MESAJLAR[3].split("<duzeltme>")[1])
            for m in MESAJLAR:
                for x in ("Ayşe", "Deneme", "Mehmet", "Zeynep", "1001", "1002", "1003"):
                    self.assertNotIn(x, m)                             # ad ve sicil sağlayıcıya gitmez
            self.assertEqual((s["kalan"], s["tur"], [s["sonuc"][n]["insan_kontrolu"] for n in ("S1", "S2", "S3")]), ([], 1, [False, True, False]))
            ws = load_workbook(cikti)["Cevaplar"]
            self.assertEqual([ws.cell(3, j).value for j in (2, 3, 8)], ["1002", "Mehmet Örnek", "Evet"])
            self.assertTrue(ws.cell(2, 7).value.startswith("Sayın Ayşe Deneme,\n\nEkim bordronuzu"))
            self.assertTrue(ws.cell(2, 7).value.endswith("Saygılarımızla,\nBordro Birimi"))
            self.assertIn("İNSAN KONTROLÜ GEREKİR", s["txt"].read_text(encoding="utf-8"))
            h = load_workbook(cikti)["Hesap"]
            self.assertIn(["S3", "1003", "Gün", 30, 22, -8], [list(r) for r in h.iter_rows(min_row=2, values_only=True)])

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_atlanan_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "s.txt"
            y.write_text("Soru: 9999 | Ekim\nmaaş?\n\nSoru: 1001 | Haziran\nmaaş?\n\nSoru: 1001 | geçen ay\nmaaş?\n", encoding="utf-8")
            s = agent.calistir(ORNEK / "bordro.csv", y, Path(tmp) / "c.xlsx", evet=True)
            self.assertEqual((len(s["hazir"]), [q["neden"][:20] for q in s["atlanan"]]),
                             (0, ["'9999' bordroda bulu", "Haziran ayına ait bo", "Ay 'geçen ay' okunam"]))
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(agent.main(["--bordro", str(Path(tmp) / "yok.csv"), "--evet"]), 1)
            y.write_text("soru yok", encoding="utf-8")
            self.assertEqual(agent.main(["--sorular", str(y), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
