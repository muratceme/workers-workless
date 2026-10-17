"""
Bordro Sorusuna Cevap Taslağı — Gerisi Bende AI Agent
İnsan Kaynakları › Bordro ve Özlük İşleri Uzmanı

Çalışanların "maaşım neden düştü", "fazla mesaim yatmamış", "bu kesinti ne" gibi sorularına bordro verisine dayanan,
anlaşılır cevap taslakları hazırlar.
1. Kod bordro dökümünü (çalışan × ay) ve soruları okur; soruyu soran çalışanın sorulan ay ve önceki ay satırlarını bulur,
   kalem kalem farkları kendisi hesaplar ve olası nedenleri işaretler: vergi dilimi değişti, gün eksik, fazla mesai / prim
   değişti, asgari ücret istisnası değişti, yeni kesinti. Ad soyadlar sağlayıcıya gönderilmez ([ÇALIŞAN]).
2. Model kodun hesapladığı rakamlarla soruyu sade bir dille cevaplar; bordrodan anlaşılamayan konularda (puantaj,
   açıklaması olmayan kesinti, hata iddiası) "insan kontrolü gerekir" der.
3. Kod cevabı denetler: cevaptaki her sayı bordro satırlarında, kodun hesapladığı farklarda ya da soruda var mı; ödeme /
   düzeltme sözü veriliyor mu; kodun işaretlediği bir konu (ör. fazla mesai 0) insan kontrolüne gönderilmiş mi.
   Sorun varsa yalnız o cevap modele geri gönderilir (en fazla --tur kez).
4. Kod hitap ("Sayın Ad Soyad,") ve imzayı ekler.
"Brüt-Net Maaş Hesaplama" kod bloğunun kişi sayfaları da bordro dökümü olarak okunabilir.

Kullanım:
    python agent.py                                              # örnek bordro ve 3 soru
    python agent.py --bordro bordro_ekim.xlsx --sorular sorular.txt --imza "Bordro Birimi"
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}
AYLAR = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


AY_KATLI = [katla(a) for a in AYLAR]


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1 or (s.count(".") == 1 and len(s.split(".")[1]) == 3):
        s = s.replace(".", "")
    try:
        return Decimal(s) if s not in ("", "-", ".") else None
    except InvalidOperation:
        return None


def tl(x: Decimal) -> str:
    return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def ay_coz(x) -> int | None:
    """'Ekim', 'Ekim 2026', '10', '2026-10', '10.2026', datetime → 10."""
    if isinstance(x, datetime):
        return x.month
    k = katla(x)
    for i, a in enumerate(AY_KATLI, 1):
        if re.search(rf"\b{a}\b", k):
            return i
    m = re.fullmatch(r"(?:\d{4} )?(\d{1,2})(?: \d{4})?", k)
    if m and 1 <= int(m[1]) <= 12:
        return int(m[1])
    return None


# ----------------------------------------------------------------------------
# Bordro dökümü
# ----------------------------------------------------------------------------

# kalem: (eş anlamlı başlıklar, görünen ad)
KALEMLER = {
    "gun": (("gun", "sgk gunu", "prim gunu", "calisma gunu"), "Gün"),
    "eksik_neden": (("eksik gun nedeni", "eksik gun aciklamasi", "eksik gun"), "Eksik gün nedeni"),
    "brut_ucret": (("brut ucret", "ciplak brut", "aylik brut"), "Brüt ücret"),
    "fazla_mesai": (("fazla mesai", "fazla calisma", "fm tutari", "fazla mesai tutari"), "Fazla mesai"),
    "prim": (("prim", "ikramiye", "ek odeme", "prim ikramiye"), "Prim / ek ödeme"),
    "brut": (("toplam brut", "brut", "brut toplam"), "Toplam brüt"),
    "sgk_isci": (("sgk isci", "sgk isci payi", "sgk primi isci"), "SGK işçi payı"),
    "issizlik_isci": (("issizlik isci", "issizlik isci payi"), "İşsizlik işçi payı"),
    "gv_matrahi": (("gv matrahi", "gelir vergisi matrahi"), "Gelir vergisi matrahı"),
    "kumulatif": (("kumulatif gv matrahi", "kumulatif matrah", "kumulatif vergi matrahi"), "Kümülatif GV matrahı"),
    "dilim": (("vergi dilimi %", "vergi dilimi", "dilim"), "Vergi dilimi (%)"),
    "hesaplanan_gv": (("hesaplanan gv", "hesaplanan gelir vergisi"), "Hesaplanan gelir vergisi"),
    "gv_istisna": (("au gv istisnasi", "asgari ucret gv istisnasi", "gv istisnasi"), "Asgari ücret GV istisnası"),
    "odenecek_gv": (("odenecek gv", "gelir vergisi", "odenecek gelir vergisi"), "Ödenecek gelir vergisi"),
    "odenecek_dv": (("odenecek dv", "damga vergisi", "odenecek damga vergisi"), "Ödenecek damga vergisi"),
    "net": (("net ucret", "net"), "Net ücret"),
    "kesinti": (("kesinti", "diger kesinti", "kesintiler", "avans kesintisi"), "Kesinti"),
    "kesinti_neden": (("kesinti aciklamasi", "kesinti nedeni", "kesinti turu"), "Kesinti açıklaması"),
    "ele_gecen": (("ele gecen", "odenecek net", "banka odemesi", "net odenen"), "Ele geçen"),
}
METIN_KALEM = {"eksik_neden", "kesinti_neden"}


def tablo_satirlari(yol: Path) -> list[tuple[str, list[list]]]:
    """(sayfa adı, satırlar) listesi; CSV tek sayfadır."""
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
        finally:
            wb.close()
    metin = yol.read_text(encoding="utf-8-sig")
    ilk = "\n".join(metin.splitlines()[:5])
    return [(yol.stem, list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count))))]


def bordro_oku(yol: Path) -> dict[str, dict]:
    """{sicil: {'ad': ..., 'aylar': {ay: {kalem: değer}}}}. Sicil sütunu yoksa Ad Soyad, o da yoksa sayfa adı kullanılır."""
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: bordro dosyası bulunamadı.")
    out: dict[str, dict] = {}
    for sayfa, satirlar in tablo_satirlari(yol):
        bas = next((i for i, r in enumerate(satirlar[:10]) if {"ay", "net ucret"} <= {katla(c) for c in r} or
                    ({"ay"} <= {katla(c) for c in r} and "net" in {katla(c) for c in r})), None)
        if bas is None:
            continue
        b = [katla(c) for c in satirlar[bas]]
        idx = {k: next((i for i, x in enumerate(b) if x in es), None) for k, (es, _) in KALEMLER.items()}
        if idx["brut"] is None and idx["brut_ucret"] is not None:
            idx["brut"], idx["brut_ucret"] = idx["brut_ucret"], None
        i_ay = b.index("ay")
        i_sicil = next((i for i, x in enumerate(b) if x in ("sicil", "sicil no", "personel no")), None)
        i_ad = next((i for i, x in enumerate(b) if x in ("ad soyad", "adi soyadi", "personel", "calisan")), None)
        for r in satirlar[bas + 1:]:
            if not r or all(c in (None, "") for c in r):
                continue
            ay = ay_coz(r[i_ay] if i_ay < len(r) else None)
            if ay is None:
                continue
            al = lambda i: None if i is None or i >= len(r) else r[i]  # noqa: E731
            ad = str(al(i_ad) or (sayfa if i_sicil is None else "")).strip()
            sicil = str(al(i_sicil) or ad).strip()
            if sicil.endswith(".0"):
                sicil = sicil[:-2]
            kayit = out.setdefault(sicil, {"ad": ad, "aylar": {}})
            kayit["ad"] = kayit["ad"] or ad
            kayit["aylar"][ay] = {k: (str(al(i) or "").strip() if k in METIN_KALEM else para(al(i))) for k, i in idx.items() if i is not None}
    if not out:
        raise llm.LLMHatasi(f"{yol.name}: 'Ay' ve 'Net Ücret' sütunları olan bordro tablosu bulunamadı.")
    return out


def sorulari_oku(yol: Path) -> list[dict]:
    """'Soru: Sicil veya Ad Soyad | Ay' satırıyla başlayan bloklar."""
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: soru dosyası bulunamadı.")
    out = []
    for s in belge.metin_oku(yol).splitlines():
        m = re.match(r"^\s*Soru\s*:\s*(.+?)\s*\|\s*(.+?)\s*$", s, re.I)
        if m:
            out.append({"no": f"S{len(out) + 1}", "kim": m[1], "ay_metin": m[2], "metin": []})
        elif out and s.strip():
            out[-1]["metin"].append(s.strip())
    if not out:
        raise llm.LLMHatasi(f"{yol.name}: 'Soru: Sicil | Ay' başlıklı soru yok.")
    for q in out:
        q["metin"] = "\n".join(q["metin"])
    return out


def calisan_bul(bordro: dict, kim: str) -> str | None:
    if kim in bordro:
        return kim
    k = katla(kim)
    adaylar = [s for s, v in bordro.items() if katla(v["ad"]) == k]
    return adaylar[0] if len(adaylar) == 1 else None


# ----------------------------------------------------------------------------
# Kodun hesabı
# ----------------------------------------------------------------------------

PARA_KALEM = [k for k in KALEMLER if k not in METIN_KALEM | {"gun", "dilim"}]


def analiz(c: dict, ay: int, soru: str) -> dict:
    """Sorulan ay ile önceki ayın kalem farkları ve kodun işaretlediği konular."""
    su = c["aylar"].get(ay)
    once = c["aylar"].get(max((a for a in c["aylar"] if a < ay), default=0))
    farklar, isaret = [], []
    if su is None:
        return {"su": None, "once": once, "farklar": [], "isaret": [("Hata", f"{AYLAR[ay - 1]} ayına ait bordro satırı yok")]}
    for k in KALEMLER:
        if k in METIN_KALEM or k not in su:
            continue
        a, b = su.get(k), (once or {}).get(k)
        if a is not None and b is not None and a != b:
            farklar.append((k, b, a, a - b))
    fk = {k: d for k, _, _, d in farklar}
    if once and "dilim" in fk:
        isaret.append(("Neden", f"Vergi dilimi %{once['dilim']:g} → %{su['dilim']:g} (kümülatif matrah {tl(su.get('kumulatif') or Decimal(0))})"))
    if once and "gun" in fk and fk["gun"] < 0:
        isaret.append(("Neden", f"Gün sayısı {once['gun']:g} → {su['gun']:g}" + (f" ({su['eksik_neden']})" if su.get("eksik_neden") else "")))
    if "gv_istisna" in fk:
        isaret.append(("Neden", f"Asgari ücret GV istisnası {tl(once['gv_istisna'])} → {tl(su['gv_istisna'])}"))
    for k in ("fazla_mesai", "prim"):
        if k in fk:
            isaret.append(("Neden", f"{KALEMLER[k][1]} {tl(once[k])} → {tl(su[k])}"))
    if (su.get("kesinti") or 0) > 0 and not su.get("kesinti_neden"):
        isaret.append(("İnsan kontrolü", f"{tl(su['kesinti'])} kesintinin açıklaması bordroda yok"))
    sk = katla(soru)
    if re.search(r"\b(fazla mesai|mesai|fazla calisma)\b", sk) and not su.get("fazla_mesai"):
        isaret.append(("İnsan kontrolü", "Soruda fazla mesai geçiyor, bu ayın bordrosunda fazla mesai yok: puantaj kontrol edilmeli"))
    if re.search(r"\b(hata|yanlis|eksik yat|fark)\b", sk) and once and not farklar:
        isaret.append(("İnsan kontrolü", "Çalışan hata bildiriyor ama iki ay arasında bordroda fark yok"))
    if not once:
        isaret.append(("Bilgi", "Önceki ay bordrosu yok; karşılaştırma yapılamadı"))
    return {"su": su, "once": once, "farklar": farklar, "isaret": isaret}


def izinli_sayilar(a: dict, soru: str, ay: int) -> set[Decimal]:
    s = {Decimal(ay), Decimal(ay - 1)}
    for satir in (a["su"] or {}), (a["once"] or {}):
        for k, v in satir.items():
            if isinstance(v, Decimal):
                s |= {v, abs(v)}
            elif isinstance(v, str):
                s |= sayilar(v)
    for _, b, c, d in a["farklar"]:
        s |= {b, c, d, abs(d)}
    return s | sayilar(soru)


def sayilar(metin: str) -> set[Decimal]:
    return {v for x in re.findall(r"\d+(?:[.,]\d+)*", str(metin)) if (v := para(x)) is not None}


def hesap_metni(a: dict, ay: int) -> str:
    on = AYLAR[ay - 2] if ay > 1 and a["once"] else "-"
    l = [f"Sorulan ay: {AYLAR[ay - 1]} · Önceki ay: {on}", "Kalem | önceki ay | sorulan ay | fark"]
    for k in KALEMLER:
        if k in METIN_KALEM or k not in (a["su"] or {}):
            continue
        b, c = (a["once"] or {}).get(k), a["su"].get(k)
        f = lambda x: "-" if x is None else (f"{x:g}" if k in ("gun", "dilim") else tl(x))  # noqa: E731
        l.append(f"{KALEMLER[k][1]} | {f(b)} | {f(c)} | {f(c - b) if b is not None and c is not None else '-'}")
    for k in METIN_KALEM:
        if (a["su"] or {}).get(k):
            l.append(f"{KALEMLER[k][1]} ({AYLAR[ay - 1]}): {a['su'][k]}")
    l.append("Kodun işaretledikleri:")
    l += [f"- [{t}] {m}" for t, m in a["isaret"]] or ["- (yok)"]
    return "\n".join(l)


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

SEMA = {"type": "object", "properties": {
    "kategori": {"type": "string", "enum": ["Vergi dilimi", "Eksik gün", "Fazla mesai / prim", "Kesinti", "Asgari ücret istisnası", "Birden fazla neden",
                                            "Bordroda fark yok", "Diğer"]},
    "cevap": S, "insan_kontrolu": {"type": "boolean"}, "kontrol_nedeni": S,
    "kullanilan_kalemler": {"type": "array", "items": S}},
    "required": ["kategori", "cevap", "insan_kontrolu", "kontrol_nedeni", "kullanilan_kalemler"], "additionalProperties": False}
SOZ = re.compile(r"\b(ödenecektir|size ödenecek|ödeyeceğiz|yatırılacak(tır)?|yatıracağız|düzeltilecek(tir)?|düzelteceğiz|iade edilecek(tir)?|"
                 r"iade edeceğiz|telafi edilecek(tir)?|garanti ediyoruz|söz veriyoruz|(bordroya|maaşınıza|hesabınıza) eklenecek(tir)?)\b", re.I)
AD_KALIBI = re.compile(r"\b(sayın|merhaba|selam)\b", re.I)


def maskele(metin: str, ad: str) -> str:
    """Soru metnindeki çalışan adını ve soyadını [ÇALIŞAN] yapar, sonra ortak maskeler."""
    for parca in sorted([ad] + ad.split(), key=len, reverse=True):
        if len(parca) > 2:
            metin = re.sub(rf"\b{re.escape(parca)}\b", "[ÇALIŞAN]", metin, flags=re.I)
    return llm.maskele(re.sub(r"(\[ÇALIŞAN\]\s*)+", "[ÇALIŞAN] ", metin).strip())


def denetle(q: dict, c: dict) -> list[str]:
    sorun = []
    fazla = sorted(sayilar(c["cevap"]) - q["izinli"], key=float)
    if fazla:
        sorun.append(f"{q['no']}: cevapta bordroda, farklarda ya da soruda olmayan sayı: {', '.join(f'{x:g}' for x in fazla)} "
                     "(yalnız <hesap> içindeki rakamları kullanın)")
    if m := SOZ.search(c["cevap"]):
        sorun.append(f"{q['no']}: cevapta ödeme / düzeltme sözü var ('{m.group(0)}'); söz vermeyin, kontrol edilip bilgi verileceğini yazın")
    zorunlu = [m for t, m in q["analiz"]["isaret"] if t in ("İnsan kontrolü", "Hata")]
    if zorunlu and not c["insan_kontrolu"]:
        sorun.append(f"{q['no']}: kodun işaretlediği konu insan kontrolü gerektiriyor ({zorunlu[0]}); insan_kontrolu true olmalı")
    if c["insan_kontrolu"] and not c["kontrol_nedeni"].strip():
        sorun.append(f"{q['no']}: insan kontrolü işaretli ama neden yazılmamış")
    if AD_KALIBI.match(c["cevap"].strip()):
        sorun.append(f"{q['no']}: cevap hitapla başlamasın; hitabı kod ekler")
    if "[ÇALIŞAN]" in c["cevap"]:
        sorun.append(f"{q['no']}: cevapta [ÇALIŞAN] yer tutucusu kalmış; adı kullanmayın")
    if len(c["cevap"]) > 1500:
        sorun.append(f"{q['no']}: cevap çok uzun ({len(c['cevap'])} karakter); 1.500 karakteri geçmesin")
    return sorun


def uret(sorular: list[dict], duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    sonuc = dict(onceki or {})
    hedef = sorular if duzeltme is None else [q for q in sorular if any(s.startswith(q["no"] + ":") for s in duzeltme)]
    for q in hedef:
        mesaj = f"<soru>\n{q['maskeli']}\n</soru>\n<hesap>\n{q['hesap']}\n</hesap>"
        d = [s.split(": ", 1)[1] for s in (duzeltme or []) if s.startswith(q["no"] + ":")]
        if d:
            mesaj += ("\n<duzeltme>\nÖnceki cevap: " + json.dumps(sonuc.get(q["no"]), ensure_ascii=False) + "\nKodun bulduğu sorunlar:\n"
                      + "\n".join(f"- {s}" for s in d) + "\n</duzeltme>")
        sonuc[q["no"]] = llm.json_iste(sistem, mesaj, SEMA)
    return sonuc


def hazirla(sorular: list[dict], bordro: dict) -> tuple[list[dict], list[dict]]:
    hazir, atlanan = [], []
    for q in sorular:
        sicil, ay = calisan_bul(bordro, q["kim"]), ay_coz(q["ay_metin"])
        if sicil is None:
            atlanan.append(q | {"neden": f"'{q['kim']}' bordroda bulunamadı (sicil ya da tam ad soyad yazın; aynı adda iki kişi olabilir)"})
            continue
        if ay is None:
            atlanan.append(q | {"neden": f"Ay '{q['ay_metin']}' okunamadı"})
            continue
        c = bordro[sicil]
        a = analiz(c, ay, q["metin"])
        if a["su"] is None:
            atlanan.append(q | {"neden": a["isaret"][0][1], "sicil": sicil, "ad": c["ad"]})
            continue
        q = q | {"sicil": sicil, "ad": c["ad"], "ay": ay, "analiz": a, "maskeli": maskele(q["metin"], c["ad"]), "hesap": hesap_metni(a, ay)}
        q["izinli"] = izinli_sayilar(a, q["maskeli"], ay)
        hazir.append(q)
    return hazir, atlanan


def tam_metin(q: dict, c: dict, imza: str) -> str:
    return f"Sayın {q['ad'] or q['sicil']},\n\n{c['cevap'].strip()}\n\nSaygılarımızla,\n{imza}"


def rapor_yaz(cikti: Path, hazir, atlanan, sonuc, kalan, tur, imza) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Cevaplar"
    _baslik(ws, ["Soru", "Sicil", "Ad Soyad", "Ay", "Soru Metni", "Kategori", "Cevap Taslağı", "İnsan Kontrolü", "Kontrol Nedeni", "Kodun İşaretledikleri",
                 "Son Kontrol"], (6, 8, 18, 8, 45, 16, 80, 9, 40, 50, 30))
    for q in hazir:
        c = sonuc.get(q["no"])
        sor = [s.split(": ", 1)[1] for s in kalan if s.startswith(q["no"] + ":")]
        isaret = "\n".join(f"[{t}] {m}" for t, m in q["analiz"]["isaret"])
        ws.append([q["no"], q["sicil"], q["ad"], AYLAR[q["ay"] - 1], q["metin"], c["kategori"], tam_metin(q, c, imza), "Evet" if c["insan_kontrolu"] else "Hayır",
                   c["kontrol_nedeni"], isaret, "; ".join(sor) or "Geçti"])
        n = ws.max_row
        for j in (6, 7, 9):
            ws.cell(n, j).fill = AI_DOLGU
        ws.cell(n, 8).fill = SARI if c["insan_kontrolu"] else YESIL
        ws.cell(n, 11).fill = KIRMIZI if sor else YESIL
    for q in atlanan:
        ws.append([q["no"], q.get("sicil", ""), q.get("ad", ""), q["ay_metin"], q["metin"], "", "Cevap hazırlanmadı: " + q["neden"], "Evet", q["neden"], "", ""])
        ws.cell(ws.max_row, 7).fill = KIRMIZI
    _hizala(ws)
    h = wb.create_sheet("Hesap")
    _baslik(h, ["Soru", "Sicil", "Kalem", "Önceki Ay", "Sorulan Ay", "Fark"], (6, 8, 28, 14, 14, 14))
    for q in hazir:
        for k, b, c, d in q["analiz"]["farklar"]:
            h.append([q["no"], q["sicil"], KALEMLER[k][1], float(b), float(c), float(d)])
            for j in (4, 5, 6):
                h.cell(h.max_row, j).number_format = "0" if k in ("gun", "dilim") else "#,##0.00"
    k = wb.create_sheet("Kontrol")
    k.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    k.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    k.cell(2, 1).fill = KIRMIZI if kalan else YESIL
    k.append(["Önemli: Cevaplar taslaktır. Göndermeden önce bordro ve puantajla karşılaştırın; 'İnsan Kontrolü: Evet' olanları "
              "kontrol etmeden göndermeyin. Rakamlar bordro dökümünden ve kodun hesapladığı farklardan alınmıştır."])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    txt = cikti.with_suffix(".txt")
    txt.write_text("\n\n".join(f"--- {q['no']} · {q['sicil']} · {AYLAR[q['ay'] - 1]}"
                               f"{' · İNSAN KONTROLÜ GEREKİR' if sonuc[q['no']]['insan_kontrolu'] else ''} ---\n{tam_metin(q, sonuc[q['no']], imza)}"
                               for q in hazir), encoding="utf-8")
    return txt


def _baslik(ws, basliklar, genislik) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font, h.alignment = BASLIK_DOLGU, BASLIK_YAZI, UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def _hizala(ws) -> None:
    for satir in ws.iter_rows(min_row=2):
        for h in satir:
            h.alignment = UST


def duzeltme_dongusu(uret_f, denetle_f, tur: int) -> tuple[dict, list[str], int]:
    taslak = uret_f(None, None)
    kalan = denetle_f(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret_f(kalan, taslak)
        kalan = denetle_f(taslak)
    return taslak, kalan, yapilan


def calistir(bordro_yolu: Path, soru_yolu: Path, cikti: Path, imza: str = "İnsan Kaynakları / Bordro Birimi", tur: int = 2,
             evet: bool = False) -> dict:
    bordro = bordro_oku(bordro_yolu)
    hazir, atlanan = hazirla(sorulari_oku(soru_yolu), bordro)
    print(f"[OK] {len(bordro)} çalışan · {len(hazir)} soru cevaplanacak · {len(atlanan)} soru atlandı")
    for q in atlanan:
        print(f"[!] {q['no']}: {q['neden']}")
    sonuc, kalan, yapilan = {}, [], 0
    if hazir:
        llm.onay_al(f"{len(hazir)} soru ve ilgili çalışanların iki aylık bordro rakamları gönderilecek. Ad soyad ve sicil gönderilmez "
                    "([ÇALIŞAN]); telefon, e-posta, TCKN, IBAN maskeli.", evet)
        sonuc, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(hazir, d, o),
                                                 lambda x: [s for q in hazir for s in denetle(q, x[q["no"]])], tur)
    txt = rapor_yaz(cikti, hazir, atlanan, sonuc, kalan, yapilan, imza)
    return {"hazir": hazir, "atlanan": atlanan, "sonuc": sonuc, "kalan": kalan, "tur": yapilan, "txt": txt}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Çalışanların bordro sorularına bordro verisine dayanan cevap taslakları hazırlar.")
    p.add_argument("--bordro", type=Path, default=ORNEK / "bordro.csv", help="Bordro dökümü (.xlsx/.csv): çalışan × ay; en az Ay ve Net Ücret sütunları")
    p.add_argument("--sorular", type=Path, default=ORNEK / "sorular.txt", help="Sorular (.txt/.docx): her soru 'Soru: Sicil veya Ad Soyad | Ay' satırıyla başlar")
    p.add_argument("--imza", default="İnsan Kaynakları / Bordro Birimi", help="Cevapların altına yazılacak imza")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "bordro_cevaplari.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.bordro, a.sorular, a.cikti, a.imza, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    ik = sum(c["insan_kontrolu"] for c in s["sonuc"].values())
    print(f"[OK] {len(s['sonuc'])} cevap taslağı · {ik} tanesi insan kontrolü gerektiriyor")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()} (+ {s['txt'].name})")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
