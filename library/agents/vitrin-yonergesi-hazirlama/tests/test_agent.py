"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

V = KLASOR / "ornek_veri"
KAMPANYA, URUNLER, VITRINLER = V / "kampanya.txt", V / "urunler.csv", V / "vitrin_tipleri.csv"
MESAJLAR = []
ADIMLAR = ["Vitrini boşaltın, camı ve zemini temizleyin.", "Ahşap dekor ve kuru yaprakları zemine yerleştirin.", "Mankenleri kombinlere göre giydirin.",
           "Vitrindeki her ürünün fiyat etiketini dışarıdan okunacak şekilde yerleştirin.", "Merkezden gelen \"Yeni Sezon\" afişini asın.",
           "Kurulumu 23.10.2026 tarihinde mağaza açılmadan tamamlayın."]


def yer(no, *sku, not_=""):
    return {"no": no, "skular": list(sku), "not": not_}


def yonerge(tip, dogru=True) -> dict:
    if tip == "Büyük":
        t = {"adimlar": list(ADIMLAR), "mankenler": [yer(1, "TR-01", "TR-03"), yer(2, "TR-02", "TR-03"), yer(3, "TR-05", "TR-07"), yer(4, "TR-06", "TR-07")],
             "podyumlar": [yer(1, "TR-04"), yer(2, "TR-02", not_="Katlanmış")], "afis_metni": "Yeni Sezon",
             "kontrol_listesi": ["Tüm fiyat etiketleri görünüyor", "Aydınlatma mankenlere yönlendirildi"]}
        if not dogru:
            t["mankenler"].append(yer(5, "TR-06"))
            t["adimlar"] = [a for a in t["adimlar"] if "fiyat" not in a]
            t["kontrol_listesi"] = ["Aydınlatma mankenlere yönlendirildi"]
    else:
        t = {"adimlar": list(ADIMLAR), "mankenler": [yer(1, "TR-01", "TR-02"), yer(2, "TR-05", "TR-06")], "podyumlar": [yer(1, "TR-04")],
             "afis_metni": "Yeni Sezon", "kontrol_listesi": ["Tüm fiyat etiketleri görünüyor"]}
        if not dogru:
            t["mankenler"] = [yer(1, "TR-01", "TR-07"), yer(2, "TR-02")]
            t["afis_metni"] = "Yeni Sezon — %40 indirim"
    return dict(t, vitrin_tipi=tip)


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return yonerge(re.search(r'<vitrin tipi="([^"]+)"', kullanici).group(1), "<duzeltme>" in kullanici)


def veri():
    return agent.vitrinleri_oku(VITRINLER), agent.urunleri_oku(URUNLER), agent.kampanyayi_oku(KAMPANYA)


class KodTesti(unittest.TestCase):
    def test_okuma(self):
        v, u, _ = veri()
        self.assertEqual([(x["tip"], x["manken"], x["podyum"], len(x["magazalar"])) for x in v], [("Büyük", 4, 2, 2), ("Küçük", 2, 1, 3)])
        self.assertEqual([x["sku"] for x in u if x["zorunlu"]], ["TR-01", "TR-02", "TR-05"])

    def test_denetim(self):
        v, u, k = veri()
        self.assertEqual(agent.denetle(v, {x["tip"]: yonerge(x["tip"]) for x in v}, u, k), [])
        s = agent.denetle(v, {x["tip"]: yonerge(x["tip"], False) for x in v}, u, k)
        self.assertEqual(len(s), 5, s)
        self.assertTrue(any(x.startswith("Büyük") and "kapasite" in x for x in s) and any(x.startswith("Büyük") and "fiyat etiketi" in x for x in s)
                        and any(x.startswith("Küçük") and "Erkek, Kadın" in x for x in s) and any(x.startswith("Küçük") and "TR-05" in x for x in s)
                        and any(x.startswith("Küçük") and "40" in x for x in s))
        # Talimat numaraları 1'den başlatmayı ister; "Manken 1 ve 2", "podyum 1'e" uydurma sayı değildir, uydurma oran yine yakalanır
        y = {x["tip"]: yonerge(x["tip"]) for x in v}
        y["Büyük"]["adimlar"] += ["Manken 1 ve 2'yi kadın, manken 3, 4 erkek kombiniyle giydirin; podyum 1'e çantayı koyun."]
        self.assertEqual(agent.denetle(v, y, u, k), [])
        y["Büyük"]["adimlar"] += ["Podyum 2'ye %30 indirim kartı koyun."]
        self.assertTrue(any("30" in x for x in agent.denetle(v, y, u, k)))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "v.xlsx"
            s = agent.calistir(KAMPANYA, URUNLER, VITRINLER, cikti, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (4, [], 1))
            self.assertNotIn("Nişantaşı", MESAJLAR[0])
            self.assertEqual(len(s["dosyalar"]), 2)
            m = s["dosyalar"][1].read_text(encoding="utf-8")
            self.assertIn("Manken 1: TR-01 Kaşe Kaban (Kahve) + TR-02 Yün Karışımlı Kazak (Hardal)", m)
            self.assertIn("Mağazalar: Bakırköy, Bursa Merkez, İzmir Alsancak", m)
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Dağıtım", "Büyük", "Küçük"])
            self.assertEqual(wb["Dağıtım"]["A6"].value, "İzmir Alsancak")

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(KAMPANYA, URUNLER, VITRINLER, Path(tmp) / "v.xlsx", tur=0, evet=True)
            self.assertEqual((len(s["kalan"]), s["dosyalar"]), (5, []))
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--evet", "--vitrinler", str(VITRINLER.with_name("yok.csv"))]), 1)


if __name__ == "__main__":
    unittest.main()
