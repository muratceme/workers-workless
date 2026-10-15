"""
Vitrin Yönergesi Hazırlama — Workers / Workless AI Agent
Görsel Düzenleme › Görsel Düzenleme Uzmanı

1. Kod kampanya bilgisini, ürün listesini (SKU; Ürün; Grup; Renk; Öncelik: Zorunlu / Opsiyonel) ve vitrin tiplerini
   (Vitrin Tipi; Manken Sayısı; Podyum Sayısı; Mağazalar) okur.
2. Model her vitrin tipi için mağazalara gönderilecek adım adım yönerge yazar: manken kombinleri, podyum / masa
   yerleşimi, afiş metni ve kontrol listesi.
3. Kod denetler:
   - Kombinlerdeki SKU'lar listede var mı; manken ve podyum sayısı vitrin tipinin kapasitesini aşıyor mu
   - "Zorunlu" her ürün her vitrin tipinde kullanılmış mı
   - Bir mankende farklı gruplardan (ör. kadın + erkek) ürün var mı
   - Yönergede fiyat etiketi adımı var mı
   - Afişte ve adımlarda kampanya bilgisinde olmayan sayı / tarih var mı
   Sorun varsa yalnız o vitrin tipinin yönergesi modele geri gönderilir (en fazla --tur kez).
4. Kod her vitrin tipi için mağazalara gönderilecek metni ve mağaza dağıtım listesini yazar.

Kullanım:
    python agent.py                                                   # örnek sonbahar vitrini
    python agent.py --kampanya kampanya.txt --urunler urunler.xlsx --vitrinler vitrin_tipleri.xlsx
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
_YER = {"type": "object", "properties": {"no": {"type": "integer"}, "skular": {"type": "array", "items": S}, "not": S},
        "required": ["no", "skular", "not"], "additionalProperties": False}
SEMA = {
    "type": "object",
    "properties": {"vitrin_tipi": S, "adimlar": {"type": "array", "items": S}, "mankenler": {"type": "array", "items": _YER},
                   "podyumlar": {"type": "array", "items": _YER}, "afis_metni": S, "kontrol_listesi": {"type": "array", "items": S}},
    "required": ["vitrin_tipi", "adimlar", "mankenler", "podyumlar", "afis_metni", "kontrol_listesi"],
    "additionalProperties": False,
}
FIYAT = re.compile(r"fiyat etiket", re.I)


def kampanyayi_oku(yol: Path) -> str:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: kampanya bilgisi bulunamadı.")
    return belge.metin_oku(yol).strip()


def urunleri_oku(yol: Path) -> list[dict]:
    r = kayitlar(yol, {"sku": ("SKU", "Stok Kodu"), "ad": ("Ürün", "Ürün Adı"), "grup": ("Grup", "Cinsiyet", "Reyon"), "renk": ("Renk",),
                       "oncelik": ("Öncelik",)}, ("sku", "ad"))
    return [{"sku": str(x["sku"]).strip(), "ad": str(x["ad"]).strip(), "grup": str(x["grup"]).strip(), "renk": str(x["renk"]).strip(),
             "zorunlu": katla(x["oncelik"]).startswith("zorunlu")} for x in r if str(x["sku"]).strip()]


def vitrinleri_oku(yol: Path) -> list[dict]:
    r = kayitlar(yol, {"tip": ("Vitrin Tipi", "Tip"), "manken": ("Manken Sayısı", "Manken"), "podyum": ("Podyum Sayısı", "Podyum", "Masa"),
                       "magazalar": ("Mağazalar", "Mağaza")}, ("tip", "manken"))
    out = []
    for x in r:
        try:
            out.append({"tip": str(x["tip"]).strip(), "manken": int(float(sayi_norm(str(x["manken"])))),
                        "podyum": int(float(sayi_norm(str(x["podyum"])))) if str(x["podyum"]).strip() else 0,
                        "magazalar": [m.strip() for m in re.split(r"[,;]", str(x["magazalar"])) if m.strip()]})
        except ValueError:
            raise llm.LLMHatasi(f"{yol.name}: '{x['tip']}' için manken / podyum sayısı okunamadı.")
    return out


def denetle(vitrinler: list[dict], sonuc: dict, urunler: list[dict], kampanya: str) -> list[str]:
    sorun, u = [], {x["sku"]: x for x in urunler}
    izinli = sayilar(kampanya)
    for v in vitrinler:
        a, t = v["tip"], sonuc.get(v["tip"])
        if not t:
            sorun.append(f"{a}: yönerge yazılmadı")
            continue
        kullanilan = set()
        for tur_, liste, kap in (("Manken", t["mankenler"], v["manken"]), ("Podyum", t["podyumlar"], v["podyum"])):
            nolar = [x["no"] for x in liste]
            if len(liste) > kap or any(not 1 <= n <= kap for n in nolar):
                sorun.append(f"{a}: {tur_.lower()} sayısı / numarası kapasiteyi aşıyor (vitrinde {kap} {tur_.lower()} var; yazılan: {nolar})")
            if len(set(nolar)) != len(nolar):
                sorun.append(f"{a}: aynı {tur_.lower()} numarası birden fazla kez kullanılmış")
            for x in liste:
                yok = [s for s in x["skular"] if s not in u]
                if yok:
                    sorun.append(f"{a}: {tur_} {x['no']}: listede olmayan SKU {', '.join(yok)}")
                kullanilan |= set(x["skular"])
                gruplar = sorted({u[s]["grup"] for s in x["skular"] if s in u and u[s]["grup"]})
                if tur_ == "Manken" and len(gruplar) > 1:
                    sorun.append(f"{a}: Manken {x['no']}: farklı gruplardan ürün ({', '.join(gruplar)})")
        eksik = [s for s, x in u.items() if x["zorunlu"] and s not in kullanilan]
        if eksik:
            sorun.append(f"{a}: zorunlu ürün kullanılmamış: {', '.join(eksik)}")
        if not any(FIYAT.search(s) for s in t["adimlar"] + t["kontrol_listesi"]):
            sorun.append(f"{a}: vitrindeki ürünlerin fiyat etiketiyle ilgili adım yok")
        for ad, m in [("Afiş", t["afis_metni"])] + [(f"Adım {i}", s) for i, s in enumerate(t["adimlar"], 1)]:
            fazla = sorted(sayilar(re.sub(r"\b[A-Z]{1,4}-?\d+\b", " ", m)) - izinli)
            if fazla:
                sorun.append(f"{a}: {ad}: kampanya bilgisinde olmayan sayı / tarih: {', '.join(fazla)}")
    return sorun


def uret(vitrinler, urunler, kampanya, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    sonuc, hedef = {}, vitrinler
    if duzeltme is not None:  # düzeltmede yalnız sorunlu vitrin tipleri yeniden yazılır
        sorunlu = {s.split(":")[0] for s in duzeltme}
        sonuc.update({a: b for a, b in (onceki or {}).items() if a not in sorunlu})
        hedef = [v for v in vitrinler if v["tip"] in sorunlu]
    liste = "\n".join(f"{x['sku']} | {x['ad']} | {x['grup'] or '-'} | {x['renk'] or '-'} | {'ZORUNLU' if x['zorunlu'] else 'opsiyonel'}" for x in urunler)
    for v in hedef:
        d = [s for s in (duzeltme or []) if s.split(":")[0] == v["tip"]]
        mesaj = (f"<kampanya>\n{kampanya}\n</kampanya>\n<urunler>\n{liste}\n</urunler>\n"
                 f"<vitrin tipi=\"{v['tip']}\" manken=\"{v['manken']}\" podyum=\"{v['podyum']}\"/>" + duzeltme_blogu(d, (onceki or {}).get(v["tip"])))
        sonuc[v["tip"]] = llm.json_iste(sistem, mesaj, SEMA) | {"vitrin_tipi": v["tip"]}
    return sonuc


def yonerge_metni(v: dict, t: dict, u: dict) -> str:
    ad = lambda s: f"{s} {u[s]['ad']} ({u[s]['renk']})" if s in u else s  # noqa: E731
    s = [f"VİTRİN YÖNERGESİ — {v['tip']} vitrin ({v['manken']} manken, {v['podyum']} podyum)", f"Mağazalar: {', '.join(v['magazalar'])}", "",
         f"Afiş: {t['afis_metni']}", "", "ADIMLAR"] + [f"{i}. {x}" for i, x in enumerate(t["adimlar"], 1)] + ["", "MANKENLER"]
    s += [f"Manken {m['no']}: " + " + ".join(ad(x) for x in m["skular"]) + (f" — {m['not']}" if m["not"] else "") for m in sorted(t["mankenler"], key=lambda m: m["no"])]
    if t["podyumlar"]:
        s += ["", "PODYUM / MASA"] + [f"Podyum {m['no']}: " + ", ".join(ad(x) for x in m["skular"]) + (f" — {m['not']}" if m["not"] else "")
                                      for m in sorted(t["podyumlar"], key=lambda m: m["no"])]
    return "\n".join(s + ["", "KONTROL LİSTESİ"] + [f"[ ] {x}" for x in t["kontrol_listesi"]] + ["", "Kurulum sonrası vitrin fotoğrafını bölge sorumlusuna gönderin."])


def rapor_yaz(cikti: Path, vitrinler, sonuc, urunler, kalan, tur) -> list[Path]:
    wb = Workbook()
    ws = wb.active
    ws.title = "Dağıtım"
    baslik(ws, ["Mağaza", "Vitrin Tipi", "Manken", "Podyum", "Yönerge Dosyası", "Kontrol"], (24, 14, 9, 9, 40, 50))
    u = {x["sku"]: x for x in urunler}
    dosyalar = []
    for v in vitrinler:
        t = sonuc.get(v["tip"])
        sor = [s.split(": ", 1)[1] for s in kalan if s.startswith(v["tip"] + ":")]
        dosya = cikti.with_name(f"{cikti.stem}_{re.sub(r'[^A-Za-z0-9]+', '_', katla(v['tip']))}.txt")
        if t and not sor:
            dosya.parent.mkdir(parents=True, exist_ok=True)
            dosya.write_text(yonerge_metni(v, t, u), encoding="utf-8")
            dosyalar.append(dosya)
        for m in v["magazalar"] or ["(mağaza yok)"]:
            ws.append([m, v["tip"], v["manken"], v["podyum"], dosya.name if t and not sor else "Gönderilmez", "; ".join(sor) or "Geçti"])
            boya(ws, [6], KIRMIZI if sor else YESIL)
        if t:
            y = wb.create_sheet(v["tip"][:31])
            baslik(y, ["Bölüm", "No", "İçerik"], (16, 6, 100))
            for i, x in enumerate(t["adimlar"], 1):
                y.append(["Adım", i, x])
            for tur_, liste in (("Manken", t["mankenler"]), ("Podyum", t["podyumlar"])):
                for m in sorted(liste, key=lambda m: m["no"]):
                    y.append([tur_, m["no"], " + ".join(f"{s} {u[s]['ad']}" if s in u else s for s in m["skular"]) + (f" — {m['not']}" if m["not"] else "")])
            y.append(["Afiş", "", t["afis_metni"]])
            for x in t["kontrol_listesi"]:
                y.append(["Kontrol", "", x])
            for r in range(2, y.max_row + 1):
                y.cell(r, 3).fill = AI_DOLGU
            hizala(y)
    son_not(ws, "Önemli: Yönergeleri mağazalara göndermeden önce görsel düzenleme sorumlusu onaylamalıdır; stok durumunu mağaza bazında kontrol edin.",
            kalan, tur)
    hizala(ws)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    return dosyalar


def calistir(kampanya_yolu: Path, urun_yolu: Path, vitrin_yolu: Path, cikti: Path, tur: int = 2, evet: bool = False) -> dict:
    kampanya = kampanyayi_oku(kampanya_yolu)
    urunler, vitrinler = urunleri_oku(urun_yolu), vitrinleri_oku(vitrin_yolu)
    if not urunler or not vitrinler:
        raise llm.LLMHatasi("Ürün listesi ya da vitrin tipi listesi boş.")
    print(f"[OK] {len(urunler)} ürün ({sum(x['zorunlu'] for x in urunler)} zorunlu) · {len(vitrinler)} vitrin tipi · "
          f"{sum(len(v['magazalar']) for v in vitrinler)} mağaza")
    llm.onay_al(f"Kampanya bilgisi, {len(urunler)} ürün ve {len(vitrinler)} vitrin tipi gönderilecek (mağaza adları gönderilmez).", evet)
    sonuc, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(vitrinler, urunler, kampanya, d, o), lambda x: denetle(vitrinler, x, urunler, kampanya), tur)
    dosyalar = rapor_yaz(cikti, vitrinler, sonuc, urunler, kalan, yapilan)
    return {"vitrinler": vitrinler, "sonuc": sonuc, "kalan": kalan, "tur": yapilan, "dosyalar": dosyalar}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    o = BURASI / "ornek_veri"
    p = argparse.ArgumentParser(description="Kampanya ve ürün listesinden mağazalara gönderilecek adım adım vitrin düzenleme yönergesi hazırlar.")
    p.add_argument("--kampanya", type=Path, default=o / "kampanya.txt", help="Kampanya bilgisi (.txt/.docx): ad, kurulum ve bitiş tarihi, tema, afiş metni")
    p.add_argument("--urunler", type=Path, default=o / "urunler.csv", help="Ürün listesi (.xlsx/.csv): SKU, Ürün; Grup, Renk, Öncelik (Zorunlu / Opsiyonel)")
    p.add_argument("--vitrinler", type=Path, default=o / "vitrin_tipleri.csv", help="Vitrin tipleri (.xlsx/.csv): Vitrin Tipi, Manken Sayısı; Podyum Sayısı, Mağazalar")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "vitrin_yonergesi.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.kampanya, a.urunler, a.vitrinler, a.cikti, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(s['dosyalar'])}/{len(s['vitrinler'])} vitrin tipi için yönerge hazır")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
