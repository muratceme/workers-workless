"""
PDF Fatura Bilgi Çıkarma — Workers / Workless AI Agent
Muhasebe › Ön Muhasebe Elemanı

E-fatura XML'i olmayan faturalar (PDF, e-postayla gelen fatura metni, .docx) için:
1. Kod belgeyi okur; IBAN'ları [IBAN-n], 11 haneli kimlik numaralarını [TCKN-n] yer tutucusuyla değiştirir (bunlar sağlayıcıya
   gönderilmez). Telefon, e-posta ve --gizle ile verilen adlar da maskelenir.
2. Model faturanın başlık bilgilerini (no, tarih, vade, satıcı, alıcı, VKN), satırlarını (mal/hizmet, miktar, birim fiyat,
   iskonto, tutar, KDV oranı ve tutarı, tevkifat) ve toplamlarını belgede yazdığı gibi çıkarır.
3. Kod çıkarımı denetler: metinler belgede aynen geçiyor mu, her sayı belgede yazıyor mu (hesaplanmış / uydurulmuş sayı
   olmasın), VKN alanı belgedeki bir numara ya da yer tutucu mu. Sorun varsa modele geri gönderilir (en fazla --tur kez).
4. Kod gerçek numaraları geri koyar ve hesap bulgularını ekler: VKN / TCKN kontrol hanesi, satır tutarı = miktar × birim
   fiyat − iskonto, KDV = tutar × oran, fatura tarihinde geçerli KDV oranı, satır ve fatura toplamları, tevkifat, mükerrer fatura.
Çıktının "Faturalar" ve "Satırlar" sayfaları "E-Fatura Okuma ve Listeleme" ile aynı biçimdedir; "Fatura KDV Tutarlılık
Kontrolü" kod bloğuna doğrudan verilebilir.

Kullanım:
    python agent.py                                   # örnek faturalarla dener
    python agent.py --girdi ./gelen_faturalar --gizle "Ayşe Deneme"
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
RENK = {"Hata": "FDE2E1", "Dikkat": "FFF4CE", "Bilgi": "E8F0FE", "Hatalı": "FDE2E1", "Kontrol": "FFF4CE", "Uygun": "E3F5E1"}
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}
TOLERANS = Decimal("0.01")
YUVARLAMA = Decimal("0.05")
YENI_ORAN_TARIHI = date(2023, 7, 10)          # KDV oranları %8 → %10, %18 → %20 (7 Temmuz 2023 tarihli Cumhurbaşkanı Kararı)
ORANLAR_YENI = {Decimal(0), Decimal(1), Decimal(10), Decimal(20)}
ORANLAR_ESKI = {Decimal(0), Decimal(1), Decimal(8), Decimal(18)}
TEVKIFAT_ORANLARI = {Decimal(n) / 10 for n in (2, 3, 4, 5, 7, 9, 10)}
AYLAR = ["ocak", "subat", "mart", "nisan", "mayis", "haziran", "temmuz", "agustos", "eylul", "ekim", "kasim", "aralik"]


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def para(x) -> Decimal | None:
    """'12.425,00' · '12,425.00' · '%20' · '5800' → Decimal. Okunamazsa None."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1 or (s.count(".") == 1 and len(s.split(".")[1]) == 3):
        s = s.replace(".", "")
    try:
        return Decimal(s) if s not in ("", "-", ".") else None
    except InvalidOperation:
        return None


def belgedeki_sayilar(metin: str) -> set[Decimal]:
    """Belgede yazan tüm sayılar (binlik ayırıcı ve ondalık virgül çözülerek)."""
    out = set()
    for x in re.findall(r"\d+(?:[.,]\d+)*", metin):
        v = para(x)
        if v is not None:
            out.add(v)
    return out


def tarih_coz(x) -> date | None:
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s[:10], f).date()
        except ValueError:
            pass
    m = re.fullmatch(r"(\d{1,2}) (\w+) (\d{4})", katla(s))
    if m and m[2] in AYLAR:
        try:
            return date(int(m[3]), AYLAR.index(m[2]) + 1, int(m[1]))
        except ValueError:
            return None
    return None


def tl(x: Decimal) -> str:
    return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def yuvarla(x: Decimal) -> Decimal:
    return x.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


# ----------------------------------------------------------------------------
# Kimlik numaraları ve maskeleme
# ----------------------------------------------------------------------------

IBAN_DESEN = re.compile(r"\bTR\s?\d{2}(?:\s?[0-9]{4}){5}\s?\d{2}\b", re.I)
TCKN_DESEN = re.compile(r"(?<!\d)[1-9]\d{10}(?!\d)")


def vkn_kontrol(no: str) -> str:
    """Boş metin = geçerli. Sağdan i. hane için c1 = (hane + i) mod 10; c1 ≠ 0 ise (c1 × 2^i) mod 9 (0 ise 9); kontrol hanesi
    = (10 − toplam mod 10) mod 10."""
    s = 0
    for i, n in enumerate(reversed(no[:9]), 1):
        c1 = (int(n) + i) % 10
        if c1:
            s += (c1 * 2 ** i) % 9 or 9
    return "" if (10 - s % 10) % 10 == int(no[9]) else "VKN kontrol hanesi tutmuyor (yazım hatası olabilir)"


def tckn_kontrol(no: str) -> str:
    d = [int(c) for c in no]
    if d[0] == 0 or d[9] != (sum(d[0:9:2]) * 7 - sum(d[1:8:2])) % 10 or d[10] != sum(d[:10]) % 10:
        return "TCKN kontrol haneleri tutmuyor (yazım hatası olabilir)"
    return ""


def kimlik_kontrol(no: str) -> str:
    if not no:
        return ""
    if re.fullmatch(r"\d{10}", no):
        return vkn_kontrol(no)
    if re.fullmatch(r"\d{11}", no):
        return tckn_kontrol(no)
    return f"'{no}' 10 haneli VKN ya da 11 haneli TCKN değil"


def yer_tut(desen: re.Pattern, onek: str, metin: str, harita: dict[str, str]) -> str:
    ters = {v: k for k, v in harita.items()}

    def degistir(m):
        deger = re.sub(r"\s", "", m.group(0)).upper()
        if deger not in ters:
            ters[deger] = f"[{onek}-{sum(k.startswith(f'[{onek}-') for k in harita) + 1}]"
            harita[ters[deger]] = deger
        return ters[deger]
    return desen.sub(degistir, metin)


def takma_adlar(terimler: list[str]) -> dict[str, str]:
    return {ad: f"[KİŞİ-{i}]" for i, ad in enumerate(dict.fromkeys(t.strip() for t in terimler if t.strip()), 1)}


def maskele(metin: str, kisiler: dict[str, str]) -> tuple[str, dict[str, str]]:
    """IBAN → [IBAN-n], 11 haneli numara → [TCKN-n], --gizle adları → [KİŞİ-n], en son ortak maskeler (telefon, e-posta)."""
    harita: dict[str, str] = {}
    metin = yer_tut(IBAN_DESEN, "IBAN", metin, harita)
    metin = yer_tut(TCKN_DESEN, "TCKN", metin, harita)
    for gercek, takma in sorted(kisiler.items(), key=lambda x: -len(x[0])):
        metin = re.sub(re.escape(gercek), takma, metin, flags=re.I)
    return llm.maskele(metin), harita


def geri_ac(s: str, kisiler: dict[str, str]) -> str:
    for gercek, takma in kisiler.items():
        s = s.replace(takma, gercek)
    return s


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

SATIR_SAYI = ["miktar", "birim_fiyat", "iskonto", "tutar", "kdv_orani", "kdv_tutari", "tevkifat_tutari"]
TOPLAM_SAYI = ["mal_hizmet_toplami", "iskonto_toplami", "toplam_kdv", "tevkifat", "vergiler_dahil", "odenecek"]
METIN = ["fatura_no", "tarih", "vade_tarihi", "satici_unvan", "alici_unvan"]
ETIKET = {"fatura_no": "Fatura no", "tarih": "Tarih", "vade_tarihi": "Vade", "satici_unvan": "Satıcı", "alici_unvan": "Alıcı",
          "satici_vkn": "Satıcı VKN/TCKN", "alici_vkn": "Alıcı VKN/TCKN", "satici_iban": "Satıcı IBAN", "urun": "Mal/hizmet",
          "miktar": "Miktar", "birim_fiyat": "Birim fiyat", "iskonto": "İskonto", "tutar": "Tutar", "kdv_orani": "KDV oranı",
          "kdv_tutari": "KDV tutarı", "tevkifat_orani": "Tevkifat oranı", "tevkifat_tutari": "Tevkif edilen KDV",
          "mal_hizmet_toplami": "Mal/hizmet toplamı", "iskonto_toplami": "İskonto toplamı", "toplam_kdv": "Toplam KDV",
          "tevkifat": "Tevkifat toplamı", "vergiler_dahil": "Vergiler dahil", "odenecek": "Ödenecek"}
SATIR = {"type": "object", "properties": {"urun": S, "birim": S, "tevkifat_orani": S, **{a: S for a in SATIR_SAYI}},
         "required": ["urun", "birim", "tevkifat_orani", *SATIR_SAYI], "additionalProperties": False}
FATURA = {"type": "object",
          "properties": {**{a: S for a in METIN + ["satici_vkn", "alici_vkn", "satici_iban"] + TOPLAM_SAYI},
                         "para_birimi": {"type": "string", "enum": ["TRY", "USD", "EUR", "GBP", "Diğer"]},
                         "satirlar": {"type": "array", "items": SATIR}},
          "required": [*METIN, "satici_vkn", "alici_vkn", "satici_iban", *TOPLAM_SAYI, "para_birimi", "satirlar"], "additionalProperties": False}
SEMA = {"type": "object", "properties": {"faturalar": {"type": "array", "items": FATURA}, "notlar": {"type": "array", "items": S}},
        "required": ["faturalar", "notlar"], "additionalProperties": False}


def denetle(c: dict, metin: str, harita: dict[str, str]) -> list[str]:
    """Modelin çıkarımındaki sorunlar (faturanın kendi hataları değil). Boş liste = sorunsuz."""
    sorun, sayilar = [], belgedeki_sayilar(metin)
    rakamlar = set(re.findall(r"(?<!\d)\d{10}(?!\d)", metin))

    def sayi(on, ad, v):
        if not v.strip():
            return
        d = para(v)
        if d is None:
            sorun.append(f"{on}: {ETIKET[ad]} '{v}' sayı olarak okunamadı; belgedeki sayıyı yazın")
        elif d not in sayilar and abs(d) not in sayilar:
            sorun.append(f"{on}: {ETIKET[ad]} {v} belgede yazmıyor (hesaplamayın; belgede yoksa boş bırakın)")

    if not c["faturalar"]:
        return []
    for i, f in enumerate(c["faturalar"], 1):
        on = f"Fatura {i}"
        for a in METIN:
            if f[a].strip() and not metinde_var(f[a], metin):
                sorun.append(f"{on}: {ETIKET[a]} '{f[a][:60]}' belgede geçmiyor (belgedeki yazımı aynen kullanın)")
        for a in ("tarih", "vade_tarihi"):
            if f[a].strip() and tarih_coz(f[a]) is None:
                sorun.append(f"{on}: {ETIKET[a]} '{f[a]}' tarih olarak okunamadı")
        for a in ("satici_vkn", "alici_vkn"):
            v = f[a].strip()
            if v and v not in harita and re.sub(r"\D", "", v) not in rakamlar:
                sorun.append(f"{on}: {ETIKET[a]} '{v}' belgede yok; belgedeki 10 haneli VKN'yi ya da [TCKN-n] yer tutucusunu yazın")
        v = f["satici_iban"].strip()
        if v and not (v.startswith("[IBAN-") and v in harita):
            sorun.append(f"{on}: Satıcı IBAN '{v[:30]}' belgedeki bir [IBAN-n] yer tutucusu değil; yalnız yer tutucuyu yazın ya da boş bırakın")
        for a in TOPLAM_SAYI:
            sayi(on, a, f[a])
        if not f["satirlar"]:
            sorun.append(f"{on}: satır yok; belgedeki mal/hizmet satırlarını yazın")
        for j, s in enumerate(f["satirlar"], 1):
            if not metinde_var(s["urun"], metin):
                sorun.append(f"{on} satır {j}: Mal/hizmet '{s['urun'][:50]}' belgede geçmiyor")
            for a in SATIR_SAYI:
                sayi(f"{on} satır {j}", a, s[a])
            if not s["tutar"].strip():
                sorun.append(f"{on} satır {j}: satır tutarı boş")
    return sorun


def bulgular(f: dict, metin: str) -> list[tuple[str, str]]:
    """Faturanın kendisiyle ilgili, kodun hesapladığı bulgular. Model sonucu değiştirmez."""
    b = []
    for a, ad in (("satici_vkn", "Satıcı"), ("alici_vkn", "Alıcı")):
        if not f[a]:
            b.append(("Hata" if a == "satici_vkn" else "Dikkat", f"{ad} VKN/TCKN belgede bulunamadı"))
        elif h := kimlik_kontrol(f[a]):
            b.append(("Hata", f"{ad} VKN/TCKN {f[a]}: {h}"))
    if not f["fatura_no"]:
        b.append(("Hata", "Fatura numarası bulunamadı"))
    t = tarih_coz(f["tarih"])
    if not t:
        b.append(("Hata", "Fatura tarihi bulunamadı"))
    gecerli = ORANLAR_YENI if not t or t >= YENI_ORAN_TARIHI else ORANLAR_ESKI
    st = sk = stv = Decimal(0)
    for j, s in enumerate(f["satirlar"], 1):
        on = f"Satır {j} ({s['urun'][:30]})"
        m, fy, isk, tut = para(s["miktar"]), para(s["birim_fiyat"]), para(s["iskonto"]) or Decimal(0), para(s["tutar"])
        oran, kdv = para(s["kdv_orani"]), para(s["kdv_tutari"])
        if tut is not None:
            st += tut
        if kdv is not None:
            sk += kdv
        if m is not None and fy is not None and tut is not None:
            fark = abs(yuvarla(m * fy - isk) - tut)
            if fark > YUVARLAMA:
                b.append(("Hata", f"{on}: {m:g} × {tl(fy)} − {tl(isk)} = {tl(yuvarla(m * fy - isk))}, faturada {tl(tut)}"))
            elif fark > TOLERANS:
                b.append(("Bilgi", f"{on}: satır tutarında {tl(fark)} yuvarlama farkı"))
        if oran is not None:
            if oran not in gecerli:
                b.append(("Hata", f"{on}: KDV oranı %{oran:g} fatura tarihinde geçerli değil (geçerli: "
                                  f"{', '.join(f'%{x:g}' for x in sorted(gecerli))})"))
            if tut is not None and kdv is not None and abs(yuvarla(tut * oran / 100) - kdv) > YUVARLAMA:
                b.append(("Hata", f"{on}: KDV {tl(tut)} × %{oran:g} = {tl(yuvarla(tut * oran / 100))}, faturada {tl(kdv)}"))
        elif tut:
            b.append(("Dikkat", f"{on}: KDV oranı belgede yok"))
        if s["tevkifat_orani"]:
            to = tevkifat_orani(s["tevkifat_orani"])
            tt = para(s["tevkifat_tutari"])
            if to not in TEVKIFAT_ORANLARI:
                b.append(("Hata", f"{on}: tevkifat oranı '{s['tevkifat_orani']}' tanımlı oranlardan değil (2, 3, 4, 5, 7, 9, 10 / 10)"))
            elif kdv is not None and tt is not None and abs(yuvarla(kdv * to) - tt) > YUVARLAMA:
                b.append(("Hata", f"{on}: tevkif edilen KDV {tl(kdv)} × {s['tevkifat_orani']} = {tl(yuvarla(kdv * to))}, faturada {tl(tt)}"))
        stv += para(s["tevkifat_tutari"]) or Decimal(0)
    mal, kdv_t, tev, dahil, oden = (para(f[a]) for a in ("mal_hizmet_toplami", "toplam_kdv", "tevkifat", "vergiler_dahil", "odenecek"))
    sinir = max(YUVARLAMA, TOLERANS * len(f["satirlar"]))

    def kiyas(ad, hesap, yazan):
        if yazan is not None and abs(hesap - yazan) > sinir:
            b.append(("Hata", f"{ad}: hesaplanan {tl(hesap)}, faturada {tl(yazan)}"))

    if f["satirlar"]:
        kiyas("Mal/hizmet toplamı (satırlar toplamı)", st, mal)
        kiyas("Toplam KDV (satır KDV'leri toplamı)", sk, kdv_t)
        if stv:
            kiyas("Tevkifat toplamı (satırlar)", stv, tev)
    if mal is not None and kdv_t is not None:
        kiyas("Vergiler dahil (mal/hizmet + KDV)", mal + kdv_t, dahil)
    if dahil is not None:
        kiyas("Ödenecek (vergiler dahil − tevkifat)", dahil - (tev or Decimal(0)), oden)
    if oden is None and dahil is None:
        b.append(("Dikkat", "Ödenecek / genel toplam belgede bulunamadı"))
    if f["para_birimi"] != "TRY":
        b.append(("Dikkat", f"Para birimi {f['para_birimi']}: kur ve TL karşılığını kontrol edin"))
    k = katla(metin[:600])
    if re.search(r"\b(kimden|from)\b", k) and re.search(r"\b(kime|to)\b", k):
        b.append(("Dikkat", "E-posta metninden okundu: e-Fatura / e-Arşiv mükellefiyetini ve faturanın aslını (GİB portalı, ETTN) kontrol edin"))
    return b


def tevkifat_orani(x) -> Decimal | None:
    """'5/10', '%50', '0,5' → 0.5"""
    s = str(x or "").strip().replace("%", "")
    m = re.fullmatch(r"(\d+)\s*/\s*(\d+)", s)
    if m:
        return Decimal(m[1]) / Decimal(m[2]) if int(m[2]) else None
    v = para(s)
    return None if v is None else v / 100 if v > 1 else v


def uret(metin: str, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    ek = ""
    if duzeltme:
        ek = ("\n<duzeltme>\nÖnceki çıkarım: " + json.dumps(onceki, ensure_ascii=False) + "\nKodun bulduğu sorunlar:\n"
              + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")
    return llm.json_iste(sistem, f"<belge>\n{metin}\n</belge>" + ek, SEMA)


def duzeltme_dongusu(uret_f, denetle_f, tur: int) -> tuple[dict, list[str], int]:
    taslak = uret_f(None, None)
    kalan = denetle_f(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret_f(kalan, taslak)
        kalan = denetle_f(taslak)
    return taslak, kalan, yapilan


def _ac(v, kisiler):
    if isinstance(v, str):
        return geri_ac(v, kisiler)
    if isinstance(v, list):
        return [_ac(x, kisiler) for x in v]
    if isinstance(v, dict):
        return {a: _ac(x, kisiler) for a, x in v.items()}
    return v


def isle(yol: Path, kisiler: dict[str, str], tur: int) -> dict:
    ham = belge.metin_oku(yol)
    if len(ham.strip()) < 20:
        raise belge.BelgeHatasi(f"{yol.name}: metin bulunamadı.")
    maskeli, harita = maskele(ham, kisiler)
    c, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(maskeli, d, o), lambda x: denetle(x, maskeli, harita), tur)
    faturalar, belgedeki = [], set(re.findall(r"(?<!\d)\d{10}(?!\d)", ham))
    for i, f in enumerate(c["faturalar"], 1):
        f = _ac(f, kisiler)
        f["satici_iban"] = harita.get(f["satici_iban"].strip(), "") if f["satici_iban"].strip().startswith("[IBAN-") else ""
        for a in ("satici_vkn", "alici_vkn"):          # belgede olmayan (uydurma) numara kullanılmaz
            v = f[a].strip()
            f[a] = harita.get(v, "") if v.startswith("[") else (re.sub(r"\D", "", v) if re.sub(r"\D", "", v) in belgedeki else "")
        f["dosya"] = yol.name
        f["bulgular"] = bulgular(f, ham)
        f["kalan"] = [s.split(": ", 1)[1] for s in kalan if re.match(rf"Fatura {i}\b", s)]
        onem = {x[0] for x in f["bulgular"]} | ({"Dikkat"} if f["kalan"] else set())
        f["durum"] = "Hatalı" if "Hata" in onem else "Kontrol" if "Dikkat" in onem else "Uygun"
        faturalar.append(f)
    return {"dosya": yol.name, "faturalar": faturalar, "notlar": [geri_ac(n, kisiler) for n in c["notlar"]], "kalan": kalan,
            "tur": yapilan, "maskelenen": len(harita)}


def mukerrer(sonuc: list[dict]) -> None:
    """Aynı satıcı + fatura no birden fazla dosyada geçiyorsa Hata ekler."""
    gorulen: dict[tuple, str] = {}
    for d in sonuc:
        for f in d["faturalar"]:
            k = (f["satici_vkn"] or katla(f["satici_unvan"]), katla(f["fatura_no"]))
            if not k[1]:
                continue
            if k in gorulen:
                f["bulgular"].append(("Hata", f"Mükerrer: aynı satıcı ve fatura no {gorulen[k]} dosyasında da var"))
                f["durum"] = "Hatalı"
            else:
                gorulen[k] = d["dosya"]


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

def _baslik(ws, basliklar, genislik) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font, h.alignment = BASLIK_DOLGU, BASLIK_YAZI, UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def _sayi(x):
    v = para(x)
    return float(v) if v is not None else None


def rapor_yaz(cikti: Path, sonuc: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Faturalar"
    _baslik(ws, ["Kaynak Dosya", "Fatura No", "Tarih", "Vade", "Para Birimi", "Satıcı VKN/TCKN", "Satıcı", "Alıcı VKN/TCKN", "Alıcı",
                 "Mal/Hizmet Toplamı", "İskonto", "Toplam KDV", "KDV Tevkifatı", "Vergiler Dahil", "Ödenecek", "Satıcı IBAN", "Satır Sayısı",
                 "Durum", "Bulgular"], (26, 20, 11, 11, 8, 14, 30, 14, 30, 14, 11, 12, 12, 14, 14, 30, 8, 10, 70))
    for d in sonuc:
        for f in d["faturalar"]:
            notlar = [f"{o}: {m}" for o, m in f["bulgular"]] + [f"Kontrol (çıkarım): {k}" for k in f["kalan"]]
            ws.append([d["dosya"], f["fatura_no"], tarih_coz(f["tarih"]) or f["tarih"], tarih_coz(f["vade_tarihi"]) or f["vade_tarihi"],
                       f["para_birimi"], f["satici_vkn"], f["satici_unvan"], f["alici_vkn"], f["alici_unvan"],
                       *(_sayi(f[a]) for a in ("mal_hizmet_toplami", "iskonto_toplami", "toplam_kdv", "tevkifat", "vergiler_dahil", "odenecek")),
                       f["satici_iban"], len(f["satirlar"]), f["durum"], "\n".join(notlar)])
            n = ws.max_row
            for j in (3, 4):
                ws.cell(n, j).number_format = "DD.MM.YYYY"
            for j in range(10, 16):
                ws.cell(n, j).number_format = "#,##0.00"
            for j in (2, 3, 4, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16):
                ws.cell(n, j).fill = AI_DOLGU
            ws.cell(n, 18).fill = PatternFill("solid", fgColor=RENK[f["durum"]])
            for h in ws[n]:
                h.alignment = UST

    s = wb.create_sheet("Satırlar")
    _baslik(s, ["Fatura No", "Tarih", "Satıcı", "Alıcı", "Sıra", "Mal/Hizmet", "Miktar", "Birim", "Birim Fiyat", "İskonto", "Tutar",
                "KDV Oranı", "KDV Tutarı", "Tevkifat Oranı", "Tevkif Edilen KDV", "Para Birimi"],
            (20, 11, 28, 28, 6, 40, 9, 8, 12, 10, 13, 9, 12, 10, 12, 8))
    for d in sonuc:
        for f in d["faturalar"]:
            for j, x in enumerate(f["satirlar"], 1):
                s.append([f["fatura_no"], tarih_coz(f["tarih"]) or f["tarih"], f["satici_unvan"], f["alici_unvan"], j, x["urun"],
                          _sayi(x["miktar"]), x["birim"], _sayi(x["birim_fiyat"]), _sayi(x["iskonto"]), _sayi(x["tutar"]),
                          _sayi(x["kdv_orani"]), _sayi(x["kdv_tutari"]), x["tevkifat_orani"], _sayi(x["tevkifat_tutari"]), f["para_birimi"]])
                s.cell(s.max_row, 2).number_format = "DD.MM.YYYY"
                for c in (9, 10, 11, 13, 15):
                    s.cell(s.max_row, c).number_format = "#,##0.00"

    b = wb.create_sheet("Bulgular")
    _baslik(b, ["Fatura No", "Satıcı", "Önem", "Bulgu", "Kaynak Dosya"], (20, 28, 9, 90, 26))
    for d in sonuc:
        for f in d["faturalar"]:
            for o, m in f["bulgular"]:
                b.append([f["fatura_no"], f["satici_unvan"], o, m, d["dosya"]])
                b.cell(b.max_row, 3).fill = PatternFill("solid", fgColor=RENK[o])
        for n in d["notlar"]:
            b.append(["—", "", "Bilgi", f"Model notu: {n}", d["dosya"]])
            b.cell(b.max_row, 3).fill = PatternFill("solid", fgColor=RENK["Bilgi"])
        if not d["faturalar"]:
            b.append(["—", "", "Dikkat", "Bu belgede fatura bulunamadı", d["dosya"]])
            b.cell(b.max_row, 3).fill = PatternFill("solid", fgColor=RENK["Dikkat"])

    k = wb.create_sheet("Kontrol")
    _baslik(k, ["Dosya", "Fatura", "Maskelenen No", "Düzeltme Turu", "Son kontrol"], (28, 8, 12, 10, 90))
    for d in sonuc:
        k.append([d["dosya"], len(d["faturalar"]), d["maskelenen"], d["tur"], "; ".join(d["kalan"]) or "Kod denetiminden sorunsuz geçti"])
        k.cell(k.max_row, 5).fill = KIRMIZI if d["kalan"] else YESIL
        k.cell(k.max_row, 5).alignment = UST
    k.append([])
    k.append([f"Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    k.append(["Önemli: Mor hücreler yapay zekâ ile çıkarılmıştır. Muhasebeleştirmeden önce faturanın aslıyla karşılaştırın. "
              "Ayrıntılı KDV ve tevkifat kontrolü için bu dosyayı 'Fatura KDV Tutarlılık Kontrolü' kod bloğuna verin."])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def dosyalar(girdi: Path) -> list[Path]:
    if girdi.is_dir():
        l = sorted(p for p in girdi.iterdir() if p.is_file() and p.suffix.lower() in belge.DESTEKLENEN)
        if not l:
            raise belge.BelgeHatasi(f"{girdi}: klasörde .pdf, .docx veya .txt dosyası yok.")
        return l
    if not girdi.exists():
        raise belge.BelgeHatasi(f"{girdi}: dosya bulunamadı.")
    return [girdi]


def calistir(girdi: Path, cikti: Path, gizle: list[str] | None = None, tur: int = 2, evet: bool = False) -> dict:
    l = dosyalar(girdi)
    kisiler = takma_adlar(gizle or [])
    llm.onay_al(f"{len(l)} fatura belgesinin metni gönderilecek. IBAN ve 11 haneli kimlik numaraları gönderilmez (yer tutucu); telefon, "
                f"e-posta{f' ve {len(kisiler)} ad' if kisiler else ''} maskeli. Unvanlar, VKN'ler, kalemler ve tutarlar gönderilir.", evet)
    sonuc = [isle(y, kisiler, tur) for y in l]
    mukerrer(sonuc)
    rapor_yaz(cikti, sonuc)
    js = cikti.with_suffix(".json")
    js.write_text(json.dumps(sonuc, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    return {"dosyalar": sonuc, "faturalar": [f for d in sonuc for f in d["faturalar"]], "json": js}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="PDF / e-posta faturalarından başlık, satır ve toplam bilgilerini çıkarır ve hesapları denetler.")
    p.add_argument("--girdi", type=Path, default=ORNEK, help="Fatura belgesi ya da klasörü (.pdf, .docx, .txt)")
    p.add_argument("--gizle", nargs="*", default=[], metavar="AD", help="Ayrıca maskelenecek kişi adları")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "faturalar.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.cikti, a.gizle, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    say = {d: sum(f["durum"] == d for f in s["faturalar"]) for d in ("Hatalı", "Kontrol", "Uygun")}
    print(f"[OK] {len(s['dosyalar'])} belge · {len(s['faturalar'])} fatura · {say['Hatalı']} hatalı, {say['Kontrol']} kontrol, {say['Uygun']} uygun")
    for f in s["faturalar"]:
        for o, m in f["bulgular"]:
            if o == "Hata":
                print(f"[!] {f['fatura_no']} ({f['satici_unvan']}): {m}")
    print(f"[OK] Rapor: {a.cikti.resolve()} (+ {s['json'].name})")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
