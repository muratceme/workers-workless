"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import json
import os
import sys
import tempfile
import unittest
from decimal import Decimal as D
from datetime import date
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
BAKIM = ORNEK / "fatura_ornek_arac_bakim.txt"
DANISMANLIK = ORNEK / "fatura_kurgu_danismanlik.txt"
CEVAP = json.loads((Path(__file__).parent / "sahte_cevaplar.json").read_text(encoding="utf-8"))
MESAJLAR = []


def bakim_hatali() -> dict:
    """Modelin sık yaptığı hatalar: belgedeki yanlış satır tutarını 'düzeltmek', vadeyi hesaplamak, unvanı kısaltmak, IBAN uydurmak."""
    c = copy.deepcopy(CEVAP["bakim"])
    f = c["faturalar"][0]
    f["satirlar"][0]["tutar"] = "3.750,00"
    f["vade_tarihi"] = "04.11.2026"
    f["satici_unvan"] = "Deneme Oto Servis A.Ş."
    f["satici_iban"] = "TR330099900001000000400001"
    return c


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    if "KURGU DANI" in kullanici:
        return copy.deepcopy(CEVAP["danismanlik"])
    if "KIRTAS" in kullanici:
        return copy.deepcopy(CEVAP["kirtasiye"])
    return copy.deepcopy(CEVAP["bakim"]) if "<duzeltme>" in kullanici else bakim_hatali()


class KodTesti(unittest.TestCase):
    def test_kimlik_ve_sayi(self):
        self.assertEqual([agent.kimlik_kontrol(x) for x in ("9876543217", "10000000146", "")], ["", "", ""])
        self.assertIn("VKN kontrol", agent.kimlik_kontrol("4567890125"))
        self.assertIn("TCKN", agent.kimlik_kontrol("10000000147"))
        self.assertIn("10 haneli", agent.kimlik_kontrol("12345"))
        self.assertEqual([agent.para(x) for x in ("12.425,00", "12,425.00", "%20", "5800", "abc")], [D("12425.00"), D("12425.00"), D(20), D(5800), None])
        self.assertEqual((agent.tarih_coz("14 Ekim 2026"), agent.tarih_coz("05.10.2026"), agent.tarih_coz("ekim ortası")), (date(2026, 10, 14), date(2026, 10, 5), None))
        self.assertEqual((agent.tevkifat_orani("5/10"), agent.tevkifat_orani("%90"), agent.tevkifat_orani("")), (D("0.5"), D("0.9"), None))

    def test_maskeleme(self):
        m, h = agent.maskele(DANISMANLIK.read_text(encoding="utf-8"), {})
        self.assertEqual(h, {"[IBAN-1]": "TR040088800005500000000123", "[TCKN-1]": "10000000146"})
        for x in ("10000000146", "TR04", "0532", "can.kurgu@"):
            self.assertNotIn(x, m)
        self.assertIn("VKN 4567890125", m)                           # 10 haneli VKN şirket bilgisidir, maskelenmez
        m, h = agent.maskele("No: KKR2026000000418 Tel 0216 555 10 20", {"Kurgu": "[KİŞİ-1]"})
        self.assertEqual((m, h), ("No: KKR2026000000418 Tel [TELEFON]", {}))

    def test_denetim(self):
        m, h = agent.maskele(BAKIM.read_text(encoding="utf-8"), {})
        self.assertEqual(agent.denetle(CEVAP["bakim"], m, h), [])
        s = agent.denetle(bakim_hatali(), m, h)
        self.assertEqual(len(s), 4)
        for parca in ("Vade '04.11.2026' belgede geçmiyor", "Satıcı 'Deneme Oto Servis A.Ş.' belgede geçmiyor", "Satıcı IBAN",
                      "Tutar 3.750,00 belgede yazmıyor"):
            self.assertTrue(any(parca in x for x in s), parca)
        c = copy.deepcopy(CEVAP["bakim"])
        c["faturalar"][0].update(alici_vkn="4567890999", satirlar=[], odenecek="yedi bin")
        s = agent.denetle(c, m, h)
        self.assertTrue(any("Alıcı VKN/TCKN '4567890999' belgede yok" in x for x in s) and any("satır yok" in x for x in s)
                        and any("sayı olarak okunamadı" in x for x in s))
        self.assertEqual(agent.denetle({"faturalar": [], "notlar": ["İrsaliye"]}, m, h), [])

    def test_bulgular(self):
        f = CEVAP["kirtasiye"]["faturalar"][0]
        self.assertEqual(agent.bulgular(f, "FATURA"), [])
        f2 = copy.deepcopy(f)
        f2["satirlar"][1].update(kdv_tutari="936,00")                 # 4.680 üzerinden KDV (iskonto düşülmemiş)
        f2.update(toplam_kdv="2.521,00", vergiler_dahil="14.900,00", odenecek="14.900,00", tarih="12.10.2022", satici_vkn="", para_birimi="EUR")
        b = agent.bulgular(f2, "FATURA")
        self.assertEqual([o for o, _ in b], ["Hata", "Hata", "Hata", "Hata", "Hata", "Hata", "Dikkat"])
        self.assertTrue(any("KDV oranı %20 fatura tarihinde geçerli değil (geçerli: %0, %1, %8, %18)" in m for _, m in b))
        self.assertTrue(any(m.startswith("Vergiler dahil") for _, m in b) and not any(m.startswith("Ödenecek") for _, m in b))
        d = agent.bulgular(agent._ac(CEVAP["danismanlik"]["faturalar"][0], {}) | {"satici_vkn": "10000000146", "alici_vkn": "4567890125"},
                           DANISMANLIK.read_text(encoding="utf-8"))
        self.assertEqual([o for o, _ in d], ["Hata", "Dikkat"])
        t = copy.deepcopy(CEVAP["danismanlik"]["faturalar"][0]) | {"satici_vkn": "10000000146", "alici_vkn": "4567890128"}
        t["satirlar"][0].update(tevkifat_orani="6/10")
        self.assertIn("tanımlı oranlardan değil", agent.bulgular(t, "")[0][1])


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ornek_duzeltme_ve_cikti(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "f.xlsx"
            s = agent.calistir(ORNEK, cikti, evet=True)
            self.assertEqual(len(MESAJLAR), 4)                      # 3 belge + bakım faturası için 1 düzeltme
            self.assertIn("Tutar 3.750,00 belgede yazmıyor", MESAJLAR[3].split("<duzeltme>")[1])
            for m in MESAJLAR:
                self.assertNotRegex(m.split("<duzeltme>")[0], r"TR\d{2}\s?\d{4}|10000000146|0532")
            self.assertEqual([(f["fatura_no"], f["durum"], f["satici_vkn"], f["satici_iban"]) for f in s["faturalar"]],
                             [("CK-2026-017", "Hatalı", "10000000146", "TR040088800005500000000123"),
                              ("KKR2026000000418", "Uygun", "9876543217", "TR480099900001000000100001"), ("DOS-2026/0931", "Hatalı", "3333333339", "")])
            self.assertEqual([len(f["bulgular"]) for f in s["faturalar"]], [2, 0, 2])
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Faturalar", "Satırlar", "Bulgular", "Kontrol"])
            ws = wb["Faturalar"]
            self.assertEqual([ws.cell(3, j).value for j in (2, 6, 10, 12, 15, 18)], ["KKR2026000000418", "9876543217", 12425, 2485, 14910, "Uygun"])
            st = wb["Satırlar"]
            self.assertEqual([c.value for c in st[1]][:13], ["Fatura No", "Tarih", "Satıcı", "Alıcı", "Sıra", "Mal/Hizmet", "Miktar", "Birim",
                                                              "Birim Fiyat", "İskonto", "Tutar", "KDV Oranı", "KDV Tutarı"])
            self.assertEqual([st.cell(2, j).value for j in (6, 11, 12, 14, 15)], ["Ekim 2026 süreç danışmanlığı hizmet bedeli", 20000, 20, "5/10", 2000])
            self.assertEqual(st.max_row, 8)

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_mukerrer_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            k = Path(tmp) / "k"
            k.mkdir()
            for ad in ("a.txt", "b.txt"):
                (k / ad).write_text((ORNEK / "fatura_kurgu_kirtasiye.txt").read_text(encoding="utf-8"), encoding="utf-8")
            s = agent.calistir(k, Path(tmp) / "f.xlsx", evet=True)
            self.assertEqual([f["durum"] for f in s["faturalar"]], ["Uygun", "Hatalı"])
            self.assertIn("Mükerrer", s["faturalar"][1]["bulgular"][0][1])
            s = agent.calistir(BAKIM, Path(tmp) / "f.xlsx", tur=0, evet=True)
            self.assertEqual((len(s["dosyalar"][0]["kalan"]), s["faturalar"][0]["satici_iban"], s["faturalar"][0]["durum"]), (4, "", "Hatalı"))
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--girdi", str(ORNEK / "yok.pdf"), "--evet"]), 1)
            bos = Path(tmp) / "bos"
            bos.mkdir()
            self.assertEqual(agent.main(["--girdi", str(bos), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
