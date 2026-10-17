"""
Parça Talebini Katalogla Eşleştirme — Gerisi Bende AI Agent
Otomotiv › Satış Sonrası › Yedek Parça Danışmanı

Atölyeden ve müşterilerden gelen serbest metin parça taleplerini ("2020 Model A ön balata 1 takım") parça kataloğuyla
eşleştirir; stok durumunu ve eksik adedi kod hesaplar.
1. Kod kataloğu ("Yedek Parça Sipariş Önerisi" kod bloğunun parça dosyası + Uyumlu Modeller, Yıl Başlangıç, Yıl Bitiş)
   ve "[GG.AA.YYYY SS:DD] Kaynak: talep" satırlarını okur. Her talep için ad benzerliğine göre aday parçaları seçer.
   Telefon, e-posta, TCKN ve --gizle adları maskelenir.
2. Model talepteki her parçayı adaylardan biriyle eşleştirir ya da belirsizse soru yazar; adet, model, yıl ve aciliyeti
   çıkarır.
3. Kod denetler: parça no adaylarda mı; istek talepte aynen var mı; adet, model ve yıl talepte yazıyor mu; seçilen parça
   modele ve yıla uyumlu mu; ön / arka, sol / sağ talepte yazanla çelişiyor mu ya da talepte hiç yazmıyor mu; model /
   yıl yazmayan talepte tahminle eşleştirme yapılmış mı; "Acil" için talepte dayanak var mı. Sorun varsa yalnız o talep
   modele geri gönderilir.
4. Kod talepleri sırayla stoktan düşer (karşılanır / kısmi / sipariş gerekir) ve kod bloğunun parça dosyasını bekleyen
   talepler eklenmiş olarak yazar.

Kullanım:
    python agent.py                                                       # örnek 5 talep
    python agent.py --talepler talepler.txt --katalog parcalar.xlsx --gizle "Ad Soyad"
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}
SATIR = re.compile(r"^\s*\[(\d{1,2}\.\d{1,2}\.\d{4})\s+(\d{1,2})[:.](\d{2})\]\s*([^:]{1,60}):\s*(.+)$")
ACIL = re.compile(r"\b(acil\w*|lifte bekli\w*|arac bekli\w*|hemen|bugun|beklemede)\b")
TARAF = (("sol", "sag"), ("on", "arka"))
BLOK_SUTUN = ["Parça No", "Parça Adı", "Stok", "Yoldaki", "Aylık Tüketim", "Bekleyen İş Emri", "Tedarik Süresi (gün)", "Paket Adedi", "Birim Fiyat"]


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def sayi(x) -> Decimal | None:
    s = re.sub(r"[^\d,.\-]", "", str(x if x is not None else ""))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1 or re.fullmatch(r"\d{1,3}\.\d{3}", s):
        s = s.replace(".", "")
    try:
        return Decimal(s) if s not in ("", "-", ".") else None
    except InvalidOperation:
        return None


def sayilar(metin: str) -> set[Decimal]:
    temiz = re.sub(r"\b[A-ZİŞĞÜÖÇ]{1,4}-[\w-]+\b", " ", str(metin))
    return {v for x in re.findall(r"\d+(?:[.,]\d+)?", temiz) if (v := sayi(x)) is not None}


def kokler(s: str) -> set[str]:
    return {w[:4] for w in katla(s).split() if len(w) >= 3 and not w.isdigit()}


def govde(ad: str) -> set[str]:
    """Taraf sözcükleri çıkarılmış ad kökleri: 'Sol Ön Far' ile 'Sağ Ön Far' aynı parçanın iki tarafıdır."""
    return kokler(ad) - {"sol", "sag", "ark"}


def taraf(s: str) -> dict[int, str]:
    k = katla(s).split()
    return {i: w for i, (a, b) in enumerate(TARAF) for w in (a, b) if w in k}


# ----------------------------------------------------------------------------
# Girdi
# ----------------------------------------------------------------------------

def katalog_oku(yol: Path) -> tuple[list[str], list[dict]]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: katalog bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, read_only=True, data_only=True)
        s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=metin.split("\n")[0].count)))
    s = [r for r in s if any(c not in (None, "") for c in r)]
    bas = [str(c or "").strip() for c in s[0]]
    kb = [katla(c) for c in bas]
    out = []
    for r in s[1:]:
        r = list(r) + [None] * (len(bas) - len(r))
        d = dict(zip(kb, r))
        no = str(d.get("parca no") or "").strip()
        if not no:
            continue
        bas_y, bit_y = sayi(d.get("yil baslangic")), sayi(d.get("yil bitis"))
        out.append({"no": no, "ad": str(d.get("parca adi") or "").strip(), "modeller": [katla(x) for x in str(d.get("uyumlu modeller") or "").split(",") if katla(x)],
                    "yil1": int(bas_y) if bas_y else None, "yil2": int(bit_y) if bit_y else None, "stok": sayi(d.get("stok")) or Decimal(0), "satir": r})
    if not out:
        raise llm.LLMHatasi(f"{yol.name}: 'Parça No' ve 'Parça Adı' sütunlarıyla en az bir parça gerekli.")
    return bas, out


def takma_adlar(adlar: list[str]) -> dict[str, str]:
    return {a: f"[KİŞİ-{i}]" for i, a in enumerate(dict.fromkeys(x.strip() for x in adlar if x.strip()), 1)}


def maskele(metin: str, kisiler: dict[str, str]) -> str:
    for g, t in sorted(kisiler.items(), key=lambda x: -len(x[0])):
        metin = re.sub(re.escape(g), t, metin, flags=re.I)
    return llm.maskele(metin)


def talepleri_oku(yol: Path, gizle: list[str] | None = None) -> list[dict]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: talep dosyası bulunamadı.")
    kisiler = takma_adlar(gizle or [])
    out = []
    for s in belge.metin_oku(yol).replace("\r", "").splitlines():
        m = SATIR.match(s)
        if m:
            out.append({"id": f"T{len(out) + 1}", "zaman": f"{m[1]} {int(m[2]):02d}:{m[3]}", "kaynak": maskele(m[4].strip(), kisiler), "metin": maskele(m[5].strip(), kisiler)})
        elif out and s.strip():
            out[-1]["metin"] += " " + maskele(s.strip(), kisiler)
    if not out:
        raise llm.LLMHatasi(f"{yol.name}: '[GG.AA.YYYY SS:DD] Kaynak: talep' biçiminde satır yok.")
    return out


def adaylar(talep: str, katalog: list[dict], n: int = 8) -> list[dict]:
    k = kokler(talep)
    puan = sorted(((len(k & kokler(p["ad"])), i, p) for i, p in enumerate(katalog)), key=lambda x: (-x[0], x[1]))
    return [p for s, _, p in puan if s > 0][:n]


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

SEMA = {"type": "object", "properties": {
    "satirlar": {"type": "array", "items": {"type": "object", "properties": {"istek": S, "parca_no": S, "adet": S, "model": S, "yil": S, "soru": S},
                                            "required": ["istek", "parca_no", "adet", "model", "yil", "soru"], "additionalProperties": False}},
    "aciliyet": {"type": "string", "enum": ["Acil", "Normal"]}},
    "required": ["satirlar", "aciliyet"], "additionalProperties": False}


def uyumlu(p: dict, model: str, yil: int | None) -> str | None:
    if p["modeller"] and katla(model) not in p["modeller"]:
        return f"{p['no']} {', '.join(p['modeller'])} için; talep {model}"
    if yil and ((p["yil1"] and yil < p["yil1"]) or (p["yil2"] and yil > p["yil2"])):
        return f"{p['no']} {p['yil1'] or '…'}–{p['yil2'] or '…'} yılları için; talep {yil}"
    return None


def denetle(t: dict, aday: list[dict], c: dict) -> list[str]:
    on, kt, sorun = t["id"], katla(t["metin"]), []
    ns = sayilar(t["metin"])
    ad = {p["no"]: p for p in aday}
    if not c["satirlar"]:
        sorun.append(f"{on}: talepte satır yok")
    for i, x in enumerate(c["satirlar"], 1):
        o = f"{on}: satır {i} ('{x['istek'][:40]}')"
        ki = katla(x["istek"])
        if not ki or ki not in kt:
            sorun.append(f"{o}: istek talepte aynen geçmiyor")
        if x["adet"].strip() and sayi(x["adet"]) not in ns:
            sorun.append(f"{o}: adet '{x['adet']}' talepte yazmıyor")
        if x["model"].strip() and katla(x["model"]) not in kt:
            sorun.append(f"{o}: model '{x['model']}' talepte yazmıyor")
        if x["yil"].strip() and x["yil"].strip() not in t["metin"]:
            sorun.append(f"{o}: yıl '{x['yil']}' talepte yazmıyor")
        if not x["parca_no"].strip():
            if not x["soru"].strip():
                sorun.append(f"{o}: parça no boş; talep sahibine sorulacak soruyu yazın")
            continue
        p = ad.get(x["parca_no"].strip())
        if not p:
            sorun.append(f"{o}: parça no '{x['parca_no']}' adaylarda yok")
            continue
        if not x["model"].strip() or (p["yil1"] or p["yil2"]) and not x["yil"].strip():
            sorun.append(f"{o}: talepte model / yıl yazmıyor; uyumluluk doğrulanamaz, parça no boş bırakılıp sorulmalı")
        elif (u := uyumlu(p, x["model"], int(x["yil"]) if x["yil"].strip().isdigit() else None)):
            sorun.append(f"{o}: uyumsuz parça: {u}")
        tp, ti = taraf(p["ad"]), taraf(t["metin"])
        for j, w in tp.items():
            if j in ti and ti[j] != w:
                sorun.append(f"{o}: talepte '{ti[j]}' yazıyor, {p['no']} '{p['ad']}'")
            elif j not in ti and any(taraf(q["ad"]).get(j) not in (None, w) and govde(q["ad"]) == govde(p["ad"]) for q in aday):
                sorun.append(f"{o}: talepte {'/'.join(TARAF[j])} belirtilmemiş, {p['no']} '{p['ad']}'; sorulmalı")
    if c["aciliyet"] == "Acil" and not ACIL.search(kt):
        sorun.append(f"{on}: talepte aciliyet ifadesi yok; 'Normal' olmalı")
    return sorun


def uret(talepler: list[dict], aday: dict, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    sonuc = dict(onceki or {})
    hedef = talepler if duzeltme is None else [t for t in talepler if any(s.startswith(t["id"] + ":") for s in duzeltme)]
    for t in hedef:
        al = "\n".join(f"{p['no']} · {p['ad']} · {', '.join(p['modeller']).title() or 'tüm modeller'} · {p['yil1'] or '…'}–{p['yil2'] or '…'}" for p in aday[t["id"]]) or "(aday yok)"
        mesaj = f"<talep zaman=\"{t['zaman']}\" kaynak=\"{t['kaynak']}\">{t['metin']}</talep>\n<adaylar>\n{al}\n</adaylar>"
        d = [s.split(": ", 1)[1] for s in (duzeltme or []) if s.startswith(t["id"] + ":")]
        if d:
            mesaj += ("\n<duzeltme>\nÖnceki cevap: " + json.dumps(sonuc.get(t["id"]), ensure_ascii=False) + "\nKodun bulduğu sorunlar:\n"
                      + "\n".join(f"- {s}" for s in d) + "\n</duzeltme>")
        sonuc[t["id"]] = llm.json_iste(sistem, mesaj, SEMA)
    return sonuc


def duzeltme_dongusu(uret_f, denetle_f, tur: int) -> tuple[dict, list[str], int]:
    taslak = uret_f(None, None)
    kalan = denetle_f(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret_f(kalan, taslak)
        kalan = denetle_f(taslak)
    return taslak, kalan, yapilan


# ----------------------------------------------------------------------------
# Hesap ve çıktı
# ----------------------------------------------------------------------------

def stok_dagit(talepler: list[dict], katalog: list[dict], sonuc: dict, kalan_sorun: list[str]) -> tuple[list[dict], dict]:
    """Talepleri sırayla stoktan düşer (acil talepler önce). Denetimden geçmeyen talep stoktan düşülmez."""
    kalan = {p["no"]: p["stok"] for p in katalog}
    toplam = defaultdict(Decimal)
    sorunlu = {s.split(":", 1)[0] for s in kalan_sorun}
    sira = sorted(talepler, key=lambda t: (sonuc[t["id"]]["aciliyet"] != "Acil", t["zaman"]))
    satirlar = []
    for t in sira:
        for x in sonuc[t["id"]]["satirlar"]:
            no, a = x["parca_no"].strip(), sayi(x["adet"])
            durum, eksik = "", None
            if not no:
                durum = "Sorulacak"
            elif t["id"] in sorunlu:
                durum = "Kontrol edin (denetimden geçmedi)"
            elif a is None:
                durum = "Adet sorulacak"
            else:
                toplam[no] += a
                v = kalan.get(no, Decimal(0))
                if v >= a:
                    durum = "Stoktan karşılanır"
                elif v > 0:
                    durum, eksik = "Kısmi", a - v
                else:
                    durum, eksik = "Sipariş gerekir", a
                kalan[no] = max(v - a, Decimal(0))
            satirlar.append({"t": t, "x": x, "durum": durum, "eksik": eksik})
    return satirlar, toplam


def tablo(ws, bas: list[str], satirlar: list[list], ai: set[int] = frozenset(), genis: dict | None = None) -> None:
    ws.append(bas)
    for h in ws[1]:
        h.fill, h.font, h.alignment = BASLIK_DOLGU, BASLIK_YAZI, UST
    for r in satirlar:
        ws.append(r)
        for j, v in enumerate(r, 1):
            ws.cell(ws.max_row, j).alignment = UST
            if j in ai and v not in (None, ""):
                ws.cell(ws.max_row, j).fill = AI_DOLGU
    for j in range(1, len(bas) + 1):
        ws.column_dimensions[get_column_letter(j)].width = (genis or {}).get(j, 14)
    ws.freeze_panes = "A2"


def n(x: Decimal | None):
    return None if x is None else (int(x) if x == x.to_integral_value() else float(x))


def rapor_yaz(cikti: Path, bas: list[str], katalog: list[dict], talepler: list[dict], sonuc: dict, dagitim: list[dict], toplam: dict, kalan: list[str], tur: int) -> Path:
    pk = {p["no"]: p for p in katalog}
    wb = Workbook()
    ws = wb.active
    ws.title = "Talepler"
    rr = [[d["t"]["id"], d["t"]["zaman"], d["t"]["kaynak"], d["t"]["metin"], sonuc[d["t"]["id"]]["aciliyet"], d["x"]["istek"], d["x"]["parca_no"],
           pk.get(d["x"]["parca_no"], {}).get("ad", ""), n(sayi(d["x"]["adet"])), d["x"]["model"], d["x"]["yil"], d["durum"], n(d["eksik"]), d["x"]["soru"],
           "; ".join(s.split(": ", 1)[1] for s in kalan if s.startswith(d["t"]["id"] + ":")) or "Geçti"] for d in dagitim]
    tablo(ws, ["Talep", "Zaman", "Kaynak", "Talep Metni", "Aciliyet", "İstek", "Parça No", "Parça Adı", "Adet", "Model", "Yıl", "Stok Durumu", "Eksik Adet",
               "Sorulacak", "Son Kontrol"], rr, ai={5, 6, 7, 9, 10, 11, 14}, genis={4: 50, 6: 30, 8: 26, 12: 22, 14: 40, 15: 30})
    renk = {"Stoktan karşılanır": YESIL, "Kısmi": SARI, "Sipariş gerekir": KIRMIZI}
    for i, r in enumerate(rr, 2):
        ws.cell(i, 12).fill = renk.get(r[11], SARI)
        ws.cell(i, 15).fill = YESIL if r[14] == "Geçti" else KIRMIZI
    g = wb.create_sheet("Parça Özeti")
    tablo(g, ["Parça No", "Parça Adı", "Stok", "Talep Edilen", "Stok Sonrası", "Eksik"],
          [[no, pk[no]["ad"], n(pk[no]["stok"]), n(a), n(max(pk[no]["stok"] - a, Decimal(0))), n(max(a - pk[no]["stok"], Decimal(0)))] for no, a in toplam.items() if no in pk],
          genis={2: 28})
    k = wb.create_sheet("Kontrol")
    k.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    k.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    k.cell(2, 1).fill = KIRMIZI if kalan else YESIL
    k.append(["Mor hücreler modelin eşleştirmesidir; stok durumu koddur (acil talepler önce düşülür). Parçayı vermeden önce şasi no ile elektronik "
              "katalogda (EPC) doğrulayın. '_parcalar.xlsx' talepler 'Bekleyen İş Emri' sütununa eklenmiş kataloğudur; 'Yedek Parça Sipariş "
              "Önerisi' kod bloğuna verilebilir."])
    k.column_dimensions["A"].width = 140
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    b = Workbook()
    kb = [katla(x) for x in bas]
    yer = {katla(c): kb.index(katla(c)) for c in BLOK_SUTUN if katla(c) in kb}
    satir = []
    for p in katalog:
        r = [p["satir"][yer[katla(c)]] if katla(c) in yer else None for c in BLOK_SUTUN]
        r[5] = n((sayi(r[5]) or Decimal(0)) + toplam.get(p["no"], Decimal(0))) if toplam.get(p["no"]) else r[5]
        satir.append(r)
    tablo(b.active, BLOK_SUTUN, satir, genis={2: 28})
    y = cikti.with_name(f"{cikti.stem}_parcalar.xlsx")
    b.save(y)
    return y


def calistir(t_yol: Path, k_yol: Path, cikti: Path, gizle: list[str] | None = None, tur: int = 2, evet: bool = False) -> dict:
    bas, katalog = katalog_oku(k_yol)
    talepler = talepleri_oku(t_yol, gizle)
    aday = {t["id"]: adaylar(t["metin"], katalog) for t in talepler}
    print(f"[OK] {len(talepler)} talep · katalogda {len(katalog)} parça")
    llm.onay_al(f"{len(talepler)} talep metni ve aday parça bilgileri (no, ad, uyumlu model, yıl) gönderilecek; stok ve fiyat gönderilmez.", evet)
    sonuc, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(talepler, aday, d, o), lambda x: [s for t in talepler for s in denetle(t, aday[t["id"]], x[t["id"]])], tur)
    dagitim, toplam = stok_dagit(talepler, katalog, sonuc, kalan)
    y = rapor_yaz(cikti, bas, katalog, talepler, sonuc, dagitim, toplam, kalan, yapilan)
    return {"talepler": talepler, "sonuc": sonuc, "dagitim": dagitim, "toplam": toplam, "kalan": kalan, "tur": yapilan, "parcalar": y}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Serbest metin parça taleplerini katalogla eşleştirir; uyumluluğu ve stok durumunu kod denetler.")
    p.add_argument("--talepler", type=Path, default=ORNEK / "talepler.txt", help="Talepler (.txt/.docx/.pdf): '[GG.AA.YYYY SS:DD] Kaynak: talep'")
    p.add_argument("--katalog", type=Path, default=ORNEK / "parcalar.csv", help="Parça No; Parça Adı; Uyumlu Modeller; Yıl Başlangıç; Yıl Bitiş; Stok; … (sipariş önerisi bloğunun sütunları)")
    p.add_argument("--gizle", nargs="*", default=[], metavar="AD", help="Taleplerde geçen ve gönderilmeyecek kişi adları")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "parca_talepleri.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    gizle = a.gizle or (["Kerem Aydın"] if a.talepler == ORNEK / "talepler.txt" else [])
    try:
        s = calistir(a.talepler, a.katalog, a.cikti, gizle, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    for d in s["dagitim"]:
        x = d["x"]
        print(f"[OK] {d['t']['id']} · {x['parca_no'] or '?'} × {x['adet'] or '?'} · {d['durum']}" + (f" (eksik {n(d['eksik'])})" if d["eksik"] else "")
              + (f" · Soru: {x['soru']}" if x["soru"] else ""))
    for x in s["kalan"]:
        print(f"[!] Kalan: {x}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[OK] Kod bloğu girdisi: {s['parcalar'].name}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
