"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import json
import os
import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
CEVAP = json.loads((Path(__file__).parent / "sahte_cevaplar.json").read_text(encoding="utf-8"))
METIN = {"ön balata": "T1", "silecek": "T2", "yağ filtresi": "T3", "far, sağ": "T4", "model yazmadı": "T5"}
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    """T4'te ilk turda 'sağ taraf' far talebini sol fara (YP-1010) eşleştirir."""
    MESAJLAR.append(kullanici)
    t = next(v for k, v in METIN.items() if k in kullanici.split("</talep>")[0])
    return copy.deepcopy(CEVAP["T4_hatali" if t == "T4" and "<duzeltme>" not in kullanici else t])


def hazir():
    bas, k = agent.katalog_oku(ORNEK / "parcalar.csv")
    t = {x["id"]: x for x in agent.talepleri_oku(ORNEK / "talepler.txt", ["Kerem Aydın"])}
    return bas, k, t


class KodTesti(unittest.TestCase):
    def test_okuma(self):
        _, k, t = hazir()
        p = {x["no"]: x for x in k}
        self.assertEqual((p["YP-1004"]["modeller"], p["YP-1004"]["yil1"], p["YP-1004"]["yil2"]), (["model a"], 2015, 2018))
        self.assertEqual(p["YP-1001"]["modeller"], ["model a", "model b"])
        self.assertEqual(t["T3"]["kaynak"], "Müşteri [KİŞİ-1] (telefon)")
        self.assertEqual([x["no"] for x in agent.adaylar(t["T4"]["metin"], k)][:1], ["YP-1011"])

    def test_denetim(self):
        _, k, t = hazir()
        d = lambda i, c: agent.denetle(t[i], agent.adaylar(t[i]["metin"], k), c)
        self.assertEqual([d(i, CEVAP[i]) for i in ("T1", "T2", "T3", "T4", "T5")], [[], [], [], [], []])
        self.assertEqual(d("T4", CEVAP["T4_hatali"]), ["T4: satır 1 ('2021 Model A far, sağ taraf, 1 adet'): talepte 'sag' yazıyor, YP-1010 'Sol Ön Far'"])
        c = copy.deepcopy(CEVAP["T2"])
        c["satirlar"][0].update(parca_no="YP-1009")                                     # 2019+ silecek, araç 2017
        self.assertIn("uyumsuz parça: YP-1009 2019–… yılları için; talep 2017", d("T2", c)[0])
        c = copy.deepcopy(CEVAP["T5"])
        c["satirlar"][0].update(parca_no="YP-1003", soru="")                            # model yokken tahmin
        s = d("T5", c)
        self.assertTrue(any("model / yıl yazmıyor" in x for x in s))
        c = copy.deepcopy(CEVAP["T1"])
        c["satirlar"][0].update(parca_no="YP-1007")                                     # arka balata
        c["satirlar"].append({"istek": "4 adet buji", "parca_no": "YP-1006", "adet": "3", "model": "Model C", "yil": "2019", "soru": ""})
        c["aciliyet"] = "Acil"
        s = d("T1", c)
        for p in ("talepte 'on' yazıyor, YP-1007", "istek talepte aynen geçmiyor", "adet '3'", "model 'Model C'", "yıl '2019'", "adaylarda yok"):
            self.assertTrue(any(p in x for x in s), p)
        c = copy.deepcopy(CEVAP["T2"]) | {"aciliyet": "Acil"}
        self.assertIn("aciliyet ifadesi yok", d("T2", c)[0])

    def test_taraf_sorusu(self):
        _, k, _ = hazir()
        t = {"id": "T9", "metin": "2021 Model A far 1 adet"}                          # sol mu sağ mı yazmıyor
        c = {"satirlar": [{"istek": "2021 Model A far 1 adet", "parca_no": "YP-1011", "adet": "1", "model": "Model A", "yil": "2021", "soru": ""}], "aciliyet": "Normal"}
        self.assertTrue(any("sol/sag belirtilmemiş" in x for x in agent.denetle(t, agent.adaylar(t["metin"], k), c)))

    def test_stok(self):
        _, k, t = hazir()
        s = {i: CEVAP[i] for i in t}
        d, top = agent.stok_dagit(list(t.values()), k, s, [])
        self.assertEqual([(x["x"]["parca_no"], x["durum"], x["eksik"]) for x in d],
                         [("YP-1003", "Stoktan karşılanır", None), ("YP-1004", "Stoktan karşılanır", None), ("YP-1001", "Adet sorulacak", None),
                          ("YP-1006", "Kısmi", D(2)), ("YP-1011", "Stoktan karşılanır", None), ("", "Sorulacak", None)])
        self.assertEqual(dict(top), {"YP-1003": D(1), "YP-1004": D(2), "YP-1006": D(4), "YP-1011": D(1)})
        d, top = agent.stok_dagit(list(t.values()), k, s, ["T1: x"])                   # denetimden geçmeyen düşülmez
        self.assertEqual(d[0]["durum"], "Kontrol edin (denetimden geçmedi)")
        self.assertNotIn("YP-1003", top)


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ornek_ve_duzeltme(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "p.xlsx"
            self.assertEqual(agent.main(["--evet", "--cikti", str(cikti)]), 0)
            self.assertEqual(len(MESAJLAR), 6)
            for m in MESAJLAR:
                self.assertNotIn("Kerem Aydın", m)
                self.assertNotIn("6.400", m)                                           # fiyat gönderilmez
            b = [[c.value for c in r] for r in load_workbook(Path(tmp) / "p_parcalar.xlsx").active.iter_rows()]
            self.assertEqual(b[0], agent.BLOK_SUTUN)
            self.assertEqual({r[0]: r[5] for r in b[1:]}["YP-1006"], 4)
            self.assertEqual({r[0]: r[5] for r in b[1:]}["YP-1001"], "5")             # adet sorulacak → eklenmez
            self.assertEqual(load_workbook(cikti)["Kontrol"]["A2"].value, "Son kontrol: Kod denetiminden sorunsuz geçti")

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(agent.main(["--katalog", str(Path(tmp) / "yok.csv"), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
