"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

DOSYA = KLASOR / "ornek_veri" / "TS-2027-014"
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return {"urun_tanimi": "[GİZLİ-1] için düşük omuzlu oversize kadın tişört.",
            "yapim_detaylari": [{"bolum": "Yaka", "detay": "1x1 ribana bisiklet yaka, 3,5 mm çift iğne çıma."},
                                {"bolum": "Baskı", "detay": "Ön göğüs, yakadan 8 cm aşağı, 22 cm genişlik."}],
            "dikis_talimatlari": ["Omuzlar overlok, atölye standardına göre."], "baski_nakis": "Su bazlı tek renk baskı.",
            "etiket_yerlesimi": ["Kompozisyon etiketinde elastan yazılmalı."], "utu_paketleme": ["Katlanıp poşetlenir."],
            "kalite_notlari": ["Boy ölçüsü L bedende kontrol edilmeli."], "bakim_talimati_onerisi": "Test sonrası belirlenecek.",
            "acik_sorular": ["XS bedeni üretilecek mi? (tasarımcı)"]}


class KontrolTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.g, cls.metinler, cls.atlanan = agent.dosyalari_oku(DOSYA)
        cls.k = agent.kontrol_et(cls.g, cls.metinler)

    def test_okuma(self):
        self.assertEqual({d.ad: d.tur for d in self.metinler}, {"malzeme_listesi.csv": "Malzeme listesi", "model_bilgisi.txt": "Model bilgisi",
                                                                "olcu_tablosu.csv": "Ölçü tablosu"})
        self.assertEqual(self.k["olcu"]["bedenler"], ["S", "M", "L", "XL"])
        self.assertEqual(self.k["model"]["bedenler"], ["XS", "S", "M", "L", "XL"])
        self.assertEqual(self.k["olcu"]["satirlar"][5]["deger"]["M"], D("18.5"))

    def test_bulgular(self):
        b = self.k["bulgular"]
        metin = "\n".join(x[1] for x in b)
        self.assertIn("ölçüsü olmayan: XS", metin)
        self.assertIn("L (69) M bedeninden (70) küçük", metin)
        self.assertIn("G Yaka genişliği (dikişten dikişe): tolerans yok", metin)
        self.assertIn("elastan", metin)
        self.assertIn("'Beden etiketi' yok", metin)
        self.assertNotIn("'Dikiş ipliği' yok", metin)                      # "ipliği" tanınır
        self.assertEqual(len(b), 5)

    def test_kompozisyon(self):
        self.assertEqual(agent.kompozisyon("%95 Pamuk %5 Likra"), [(D(95), "pamuk"), (D(5), "likra")])
        self.assertEqual(agent.kompozisyon("60% cotton, 40% polyester"), [(D(60), "cotton"), (D(40), "polyester")])
        self.assertEqual(agent.kompozisyon("Pamuk %100"), [(D(100), "pamuk")])

    def test_kompozisyon_toplami_ve_duzensiz_artis(self):
        g = dict(self.g)
        g["bom"] = [dict(self.g["bom"][0], komp="%90 Pamuk %5 Elastan")]
        g["olcu"] = {"bedenler": ["XS", "S", "M", "L", "XL"],
                     "satirlar": [{"kod": "A", "ad": "Göğüs", "tol": D(1), "deger": dict(zip(["XS", "S", "M", "L", "XL"], map(D, (50, 52, 54, 62, 64))))}]}
        metin = "\n".join(x[1] for x in agent.kontrol_et(g, self.metinler)["bulgular"])
        self.assertIn("kompozisyon toplamı %95", metin)
        self.assertIn("M→L artışı 8 cm", metin)

    def test_beden_araligi(self):
        self.assertEqual(agent.beden_araligi("XS-XL"), ["XS", "S", "M", "L", "XL"])
        self.assertEqual(agent.beden_araligi("36-42"), ["36", "38", "40", "42"])
        self.assertEqual(agent.beden_araligi("S, M, L"), ["S", "M", "L"])


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_uctan_uca(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "f.xlsx"
            s = agent.calistir(DOSYA, cikti, ["Kurgu Moda"], evet=True)
            m = MESAJLAR[0]
            self.assertNotIn("Kurgu Moda", m)
            for var in ("<tablolar>", "B | Boy", "%95 Pamuk %5 Likra", "<kod_kontrolleri>"):
                self.assertIn(var, m)
            self.assertEqual(s["yanit"]["dogrulanamayan_sayilar"], ["3,5 mm"])     # 8 cm ve 22 cm tasarım notunda var
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Kapak", "Ölçü Tablosu", "Malzeme Listesi", "Yapım Detayları", "Talimatlar", "Kontroller",
                                             "Revizyonlar"])
            ol = wb["Ölçü Tablosu"]
            self.assertEqual([c.value for c in ol[1]], ["Kod", "Ölçü Noktası", "Tolerans ±", "S", "M", "L", "XL"])
            self.assertEqual(ol.cell(3, 6).value, 69)                               # ölçüler aynen aktarılır (hata dahil)
            self.assertTrue(ol.cell(2, 5).font.bold)                                # ana beden M
            self.assertIn("Kurgu Moda", wb["Kapak"]["B5"].value + wb["Kapak"]["B9"].value)
            md = s["md"].read_text(encoding="utf-8")
            self.assertIn("Kurgu Moda için", md)
            self.assertIn("3,5 mm", md)

    def test_klasor_yok(self):
        with self.assertRaises(llm.LLMHatasi):
            agent.calistir(Path("olmayan_klasor"), Path("x.xlsx"), evet=True)


if __name__ == "__main__":
    unittest.main()
