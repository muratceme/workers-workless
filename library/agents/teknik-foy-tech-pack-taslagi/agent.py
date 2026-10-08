"""
Teknik Föy (Tech Pack) Taslağı — Workers / Workless AI Agent
Tekstil ve Hazır Giyim › Tasarım ve Koleksiyon › Moda Tasarımcısı

1. Model klasöründeki dosyalar okunur: model bilgisi ve tasarım notu (.txt/.docx/.pdf), ölçü tablosu
   (Kod, Ölçü Noktası, Tolerans, bedenler) ve malzeme listesi (BOM: kumaş, aksesuar, etiket, ambalaj).
2. Kod kontrol eder: bedenler arası ölçü sırası ve düzensiz artışlar, eksik tolerans, model bilgisindeki beden
   aralığı ↔ ölçü tablosu, kumaşta kompozisyon ve gramaj, kompozisyon toplamı %100, lif adları (ör. "likra" marka
   adıdır → elastan), zorunlu etiket/aksesuar kalemleri, renk listesi.
3. Model; ürün tanımı, yapım detayları, dikiş talimatları, etiket yerleşimi, ütü/paketleme ve kalite notlarını
   taslak olarak yazar. Ölçüler ve malzemeler koddan gelir; modelin metnindeki girdilerde olmayan sayılar işaretlenir.
4. Çıktı: Excel teknik föy (Kapak + onay alanları, Ölçü Tablosu, Malzeme Listesi, Yapım Detayları, Talimatlar,
   Kontroller, Revizyonlar) + Markdown özet.

Kullanım:
    python agent.py                                          # örnek modelle dener
    python agent.py --girdi ./TS-2027-014 --gizle "Müşteri Markası"
"""
from __future__ import annotations

import argparse
import csv
import re
import statistics
import sys
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent

BEDEN_SIRASI = ["3XS", "XXS", "2XS", "XS", "S", "M", "L", "XL", "XXL", "2XL", "XXXL", "3XL", "4XL", "5XL"]
# Lif adı yazımları: (aranan, doğru ad, açıklama). Etikette ticari marka değil genel lif adı kullanılır (AB 1007/2011).
LIF_UYARI = [("likra", "elastan", "Lycra/likra ticari marka adıdır; etikette lif adı 'elastan' yazılmalı"),
             ("lycra", "elastan", "Lycra ticari marka adıdır; etikette lif adı 'elastan' yazılmalı"),
             ("spandex", "elastan", "'spandex' Türkçe etikette 'elastan' olarak yazılmalı"),
             ("naylon", "poliamid", "'naylon' yerine lif adı 'poliamid' yazılmalı"),
             ("tencel", "lyocell", "Tencel ticari marka adıdır; lif adı 'lyocell' yazılmalı"),
             ("bambu", "viskon", "bambu tek başına lif adı değildir; çoğunlukla 'viskon (bambu kaynaklı)' gibi yazılır — tedarikçi belgesiyle doğrulayın"),
             ("polyester", None, None), ("pamuk", None, None)]
LIFLER = ("pamuk", "polyester", "elastan", "viskon", "modal", "lyocell", "poliamid", "yün", "keten", "akrilik", "ipek", "kaşmir",
          "polipropilen", "asetat", "kenevir", "metalik", "likra", "lycra", "spandex", "naylon", "tencel", "bambu", "organik pamuk",
          "geri dönüştürülmüş polyester", "geri dönüştürülmüş pamuk",
          "cotton", "elastane", "viscose", "wool", "linen", "polyamide", "acrylic", "silk")
ZORUNLU = [("Dikiş ipliği", ("iplik", "ipliğ")), ("Marka etiketi", ("marka etiket", "ana etiket", "boyun etiket")),
           ("Beden etiketi", ("beden etiket",)), ("Yıkama / kompozisyon etiketi", ("yıkama", "kompozisyon", "bakım etiket")),
           ("Poşet / ambalaj", ("poşet", "polibag", "ambalaj"))]


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def katla(s) -> str:
    s = kucuk(s)
    for a, b in zip("ıöüşğç", "iousgc"):
        s = s.replace(a, b)
    return re.sub(r"[^a-z0-9%]+", " ", s).strip()


def sayi(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace(" ", "").replace(",", ".")
    s = s.replace("±", "").replace("+/-", "").replace("cm", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def sayi_yaz(x) -> str:
    return "—" if x is None else f"{x.normalize():f}".replace(".", ",")


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [[("" if c is None else str(c).strip()) for c in r] for r in satirlar if any(c not in (None, "") for c in r)]


# ----------------------------------------------------------------------------
# Girdiler
# ----------------------------------------------------------------------------

@dataclass
class Dosya:
    ad: str
    tur: str
    metin: str


def olcu_oku(yol: Path) -> dict:
    s = tablo_oku(yol)
    bi = next((i for i, r in enumerate(s[:10]) if any(katla(c) in ("olcu noktasi", "pom", "olcu") for c in r)), 0)
    b = s[bi]
    kb = [katla(c) for c in b]
    i_kod = next((kb.index(a) for a in ("kod", "pom kodu", "no") if a in kb), None)
    i_ad = next(kb.index(a) for a in ("olcu noktasi", "pom", "olcu") if a in kb)
    i_tol = next((kb.index(a) for a in ("tolerans", "tol", "tolerans cm") if a in kb), None)
    bedenler = [(j, b[j].strip().upper()) for j in range(len(b)) if b[j].strip().upper() in BEDEN_SIRASI or re.fullmatch(r"\d{2,3}", b[j].strip())]
    satirlar = []
    for r in s[bi + 1:]:
        if len(r) <= i_ad or not r[i_ad]:
            continue
        satirlar.append({"kod": r[i_kod] if i_kod is not None and i_kod < len(r) else "", "ad": r[i_ad],
                         "tol": sayi(r[i_tol]) if i_tol is not None and i_tol < len(r) else None,
                         "deger": {bd: sayi(r[j]) if j < len(r) else None for j, bd in bedenler}})
    return {"bedenler": [bd for _, bd in bedenler], "satirlar": satirlar}


def bom_oku(yol: Path) -> list[dict]:
    s = tablo_oku(yol)
    bi = next((i for i, r in enumerate(s[:10]) if any(katla(c) in ("malzeme", "malzeme adi", "tur") for c in r)), 0)
    kb = [katla(c) for c in s[bi]]

    def i(*adlar):
        return next((kb.index(a) for a in adlar if a in kb), None)
    idx = {"tur": i("tur", "malzeme turu"), "kod": i("malzeme kodu", "kod"), "ad": i("malzeme", "malzeme adi", "aciklama"),
           "komp": i("kompozisyon", "icerik", "lif icerigi"), "gramaj": i("gramaj", "gramaj g m2", "gramaj g m²"),
           "en": i("en", "en cm"), "renk": i("renk", "renkler"), "tedarikci": i("tedarikci",), "yer": i("kullanim yeri", "yer", "kullanildigi yer"),
           "tuketim": i("tuketim", "birim tuketim"), "birim": i("birim",)}
    al = lambda r, k: r[idx[k]] if idx[k] is not None and idx[k] < len(r) else ""  # noqa: E731
    return [{k: al(r, k) for k in idx} for r in s[bi + 1:] if al(r, "ad")]


def dosyalari_oku(klasor: Path) -> tuple[dict, list[Dosya], list[str]]:
    g = {"olcu": None, "bom": None, "model": None}
    metinler, atlanan = [], []
    for yol in sorted(klasor.iterdir()):
        ad = katla(yol.stem)
        if yol.suffix.lower() in {".csv", ".xlsx", ".xlsm"}:
            ilk = katla(" ".join(c for r in tablo_oku(yol)[:4] for c in r))
            if "olcu" in ad or "olcu noktasi" in ilk:
                g["olcu"], tur = olcu_oku(yol), "Ölçü tablosu"
            elif "malzeme" in ad or "bom" in ad or "kompozisyon" in ilk:
                g["bom"], tur = bom_oku(yol), "Malzeme listesi"
            else:
                atlanan.append(f"{yol.name} (tanınmayan tablo)")
                continue
            metinler.append(Dosya(yol.name, tur, ""))
            continue
        if yol.suffix.lower() not in belge.DESTEKLENEN:
            atlanan.append(yol.name)                       # çizim/görseller okunmaz; föyde referans olarak listelenir
            continue
        try:
            metin = belge.metin_oku(yol)
        except belge.BelgeHatasi as h:
            atlanan.append(f"{yol.name} ({h})")
            continue
        tur = "Model bilgisi" if ("model" in ad or "bilgi" in ad) and g["model"] is None else "Tasarım / müşteri notu"
        if tur == "Model bilgisi":
            g["model"] = metin
        metinler.append(Dosya(yol.name, tur, metin))
    return g, metinler, atlanan


def alan(metin: str | None, *etiketler: str) -> str | None:
    for et in etiketler:
        m = re.search(r"^\s*" + re.escape(et) + r"\s*:\s*(.+)$", metin or "", re.I | re.M)
        if m:
            return m.group(1).strip()
    return None


def beden_araligi(s: str | None) -> list[str]:
    if not s:
        return []
    parca = [p.strip().upper() for p in re.split(r"[,;/ ]+", s) if p.strip()]
    if len(parca) == 1 and "-" in parca[0]:
        a, b = parca[0].split("-", 1)
        if a in BEDEN_SIRASI and b in BEDEN_SIRASI:
            sira = [x for x in BEDEN_SIRASI if x not in ("2XS", "2XL", "3XL") or x in (a, b)]
            return sira[sira.index(a):sira.index(b) + 1]
        if a.isdigit() and b.isdigit():
            return [str(x) for x in range(int(a), int(b) + 1, 2)]
    return parca


def model_bilgisi(metin: str | None) -> dict:
    return {"kod": alan(metin, "Model Kodu", "Model No", "Style No"), "ad": alan(metin, "Model Adı", "Ürün Adı", "Style Name"),
            "sezon": alan(metin, "Sezon", "Season"), "musteri": alan(metin, "Müşteri", "Marka", "Müşteri / Marka"),
            "bedenler": beden_araligi(alan(metin, "Beden Aralığı", "Bedenler", "Size Range")),
            "renkler": [x.strip() for x in re.split(r"[,;]", alan(metin, "Renkler", "Renk", "Colorways") or "") if x.strip()],
            "ana_beden": (alan(metin, "Ana Beden", "Numune Bedeni", "Base Size") or "").upper() or None}


# ----------------------------------------------------------------------------
# Kod kontrolleri
# ----------------------------------------------------------------------------

def kompozisyon(s: str) -> list[tuple[Decimal, str]]:
    """'%95 Pamuk %5 Elastan' / '95% pamuk, 5% elastan' / 'Pamuk %100' → [(95, 'pamuk'), (5, 'elastan')]."""
    k = kucuk(s)
    sonuc = []
    for m in re.finditer(r"%\s*(\d+(?:[.,]\d+)?)\s*([a-zçğıöşü ]+?)(?=%|\d|,|;|$)|(\d+(?:[.,]\d+)?)\s*%\s*([a-zçğıöşü ]+?)(?=%|\d|,|;|$)", k):
        oran_, lif = (m.group(1), m.group(2)) if m.group(1) else (m.group(3), m.group(4))
        sonuc.append((Decimal(oran_.replace(",", ".")), lif.strip()))
    if not sonuc:
        m = re.search(r"([a-zçğıöşü ]+?)\s*%\s*(\d+)", k)
        if m:
            sonuc.append((Decimal(m.group(2)), m.group(1).strip()))
    return sonuc


def kontrol_et(g: dict, metinler: list[Dosya]) -> dict:
    bulgular = []

    def b(onem, aciklama, kaynak=""):
        bulgular.append((onem, aciklama, kaynak))

    mb = model_bilgisi(g["model"])
    if g["model"] is None:
        b("yüksek", "Model bilgisi dosyası yok (adında 'model' geçen .txt/.docx: Model Kodu, Model Adı, Sezon, Beden Aralığı, Renkler…)")
    else:
        for ad, k in (("Model Kodu", "kod"), ("Model Adı", "ad"), ("Sezon", "sezon"), ("Beden Aralığı", "bedenler"), ("Renkler", "renkler")):
            if not mb[k]:
                b("orta", f"model bilgisinde '{ad}' yok", "model bilgisi")

    o = g["olcu"]
    if not o:
        b("yüksek", "Ölçü tablosu yok")
    else:
        bd = o["bedenler"]
        if mb["bedenler"] and bd and mb["bedenler"] != bd:
            eksik = [x for x in mb["bedenler"] if x not in bd]
            fazla = [x for x in bd if x not in mb["bedenler"]]
            b("yüksek", f"beden aralığı ({', '.join(mb['bedenler'])}) ölçü tablosuyla ({', '.join(bd)}) uyuşmuyor"
                        + (f"; ölçüsü olmayan: {', '.join(eksik)}" if eksik else "") + (f"; fazladan: {', '.join(fazla)}" if fazla else ""),
              "model bilgisi, ölçü tablosu")
        if mb["ana_beden"] and mb["ana_beden"] not in bd:
            b("orta", f"ana (numune) beden {mb['ana_beden']} ölçü tablosunda yok", "ölçü tablosu")
        for r in o["satirlar"]:
            ref = f"{r['kod']} {r['ad']}".strip()
            if r["tol"] is None:
                b("orta", f"{ref}: tolerans yok", "ölçü tablosu")
            degerler = [(x, r["deger"].get(x)) for x in bd]
            bos = [x for x, v in degerler if v is None]
            if bos:
                b("orta", f"{ref}: {', '.join(bos)} bedeninde ölçü yok", "ölçü tablosu")
            dolu = [(x, v) for x, v in degerler if v is not None]
            adimlar = [(dolu[i][0], dolu[i + 1][0], dolu[i + 1][1] - dolu[i][1]) for i in range(len(dolu) - 1)]
            for a, c, f in adimlar:
                if f < 0:
                    b("yüksek", f"{ref}: {c} ({sayi_yaz(dict(dolu)[c])}) {a} bedeninden ({sayi_yaz(dict(dolu)[a])}) küçük — yazım hatası?",
                      "ölçü tablosu")
            pozitif = [f for _, _, f in adimlar if f > 0]
            if len(pozitif) >= 3:
                med = Decimal(str(statistics.median(pozitif)))
                for a, c, f in adimlar:
                    if f > 0 and med > 0 and f > med * Decimal("2.5"):
                        b("orta", f"{ref}: {a}→{c} artışı {sayi_yaz(f)} cm, diğer adımların medyanı {sayi_yaz(med)} cm — düzensiz artış",
                          "ölçü tablosu")
            if r["tol"] is not None and adimlar and any(0 < f < r["tol"] for _, _, f in adimlar):
                b("bilgi", f"{ref}: beden artışı toleranstan küçük (tolerans ±{sayi_yaz(r['tol'])} cm), ardışık bedenler ölçüyle ayırt edilemeyebilir",
                  "ölçü tablosu")

    bom = g["bom"] or []
    if not bom:
        b("yüksek", "Malzeme listesi (BOM) yok")
    tum = " ".join(kucuk(f"{x['tur']} {x['ad']}") for x in bom)
    for x in bom:
        ref = x["ad"]
        kumas = "kumaş" in kucuk(x["tur"]) or "kumas" in katla(x["tur"])
        if kumas:
            if not x["komp"]:
                b("yüksek", f"{ref}: kumaşın kompozisyonu yok (etiket için zorunlu)", "malzeme listesi")
            if not x["gramaj"]:
                b("orta", f"{ref}: gramaj (g/m²) yok", "malzeme listesi")
        if x["komp"]:
            k = kompozisyon(x["komp"])
            if not k:
                b("orta", f"{ref}: kompozisyon okunamadı ('{x['komp']}'); '%95 pamuk %5 elastan' biçiminde yazın", "malzeme listesi")
            else:
                toplam = sum(o_ for o_, _ in k)
                if toplam != 100:
                    b("yüksek", f"{ref}: kompozisyon toplamı %{sayi_yaz(toplam)} (≠ %100): '{x['komp']}'", "malzeme listesi")
                for _, lif in k:
                    for aranan, dogru, aciklama in LIF_UYARI:
                        if aranan in lif and aciklama:
                            b("orta", f"{ref}: {aciklama}", "malzeme listesi")
                    if not any(l_ in lif for l_ in LIFLER):
                        b("bilgi", f"{ref}: '{lif}' tanınan lif adları arasında değil; yazımı kontrol edin", "malzeme listesi")
    for ad, kelimeler in ZORUNLU:
        if bom and not any(k in tum for k in kelimeler):
            b("orta", f"malzeme listesinde '{ad}' yok", "malzeme listesi")
    renk_bom = {x["renk"].strip() for x in bom if "kuma" in katla(x["tur"]) and x["renk"].strip() and katla(x["renk"]) not in ("renge gore", "ana renk", "tum renkler", "ton ton")}
    for r in sorted(renk_bom):
        if mb["renkler"] and katla(r) not in {katla(x) for x in mb["renkler"]} and not any(katla(x) in katla(r) for x in mb["renkler"]):
            b("bilgi", f"malzeme listesindeki renk '{r}' model renkleri ({', '.join(mb['renkler'])}) arasında yok", "malzeme listesi")
    return {"bulgular": bulgular, "model": mb, "olcu": o, "bom": bom}


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

SEMA = {
    "type": "object",
    "properties": {
        "urun_tanimi": {"type": "string"},
        "yapim_detaylari": {"type": "array", "items": {
            "type": "object", "properties": {"bolum": {"type": "string"}, "detay": {"type": "string"}},
            "required": ["bolum", "detay"], "additionalProperties": False}},
        "dikis_talimatlari": {"type": "array", "items": {"type": "string"}},
        "baski_nakis": {"type": "string"},
        "etiket_yerlesimi": {"type": "array", "items": {"type": "string"}},
        "utu_paketleme": {"type": "array", "items": {"type": "string"}},
        "kalite_notlari": {"type": "array", "items": {"type": "string"}},
        "bakim_talimati_onerisi": {"type": "string"},
        "acik_sorular": {"type": "array", "items": {"type": "string"}},
    },
    "required": ["urun_tanimi", "yapim_detaylari", "dikis_talimatlari", "baski_nakis", "etiket_yerlesimi", "utu_paketleme",
                 "kalite_notlari", "bakim_talimati_onerisi", "acik_sorular"],
    "additionalProperties": False,
}


def gizle(metin: str, terimler: list[str]) -> str:
    for i, t in enumerate(sorted({t.strip() for t in terimler if t.strip()}, key=len, reverse=True), 1):
        metin = re.sub(re.escape(t), f"[GİZLİ-{i}]", metin, flags=re.I)
    return llm.maskele(metin)


def geri_ac(metin: str, terimler: list[str]) -> str:
    for i, t in enumerate(sorted({t.strip() for t in terimler if t.strip()}, key=len, reverse=True), 1):
        metin = metin.replace(f"[GİZLİ-{i}]", t)
    return metin


def tablo_metni(k: dict) -> list[str]:
    s = []
    if k["olcu"]:
        s.append("Ölçü tablosu (cm): Kod | Ölçü noktası | Tolerans | " + " | ".join(k["olcu"]["bedenler"]))
        for r in k["olcu"]["satirlar"]:
            s.append(f"{r['kod']} | {r['ad']} | ±{sayi_yaz(r['tol'])} | " + " | ".join(sayi_yaz(r["deger"].get(x)) for x in k["olcu"]["bedenler"]))
    if k["bom"]:
        s.append("Malzeme listesi: Tür | Malzeme | Kompozisyon | Gramaj | Renk | Kullanım yeri | Tüketim")
        for x in k["bom"]:
            s.append(" | ".join(x[c] or "—" for c in ("tur", "ad", "komp", "gramaj", "renk", "yer", "tuketim")))
    return s


def sayilar(metin: str) -> set[str]:
    return {m.replace(",", ".").rstrip("0").rstrip(".") if "." in m.replace(",", ".") else m
            for m in re.findall(r"\d+(?:[.,]\d+)?", metin)}


def model_yaz(metinler: list[Dosya], k: dict, terimler: list[str]) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    kk = [f"- [{o}] {a}" for o, a, _ in k["bulgular"]]
    mesaj = gizle("\n".join([
        "<belgeler>", *[f'<belge ad="{d.ad}" tur="{d.tur}">\n{d.metin}\n</belge>' for d in metinler if d.metin], "</belgeler>",
        "<tablolar>", *tablo_metni(k), "</tablolar>",
        "<kod_kontrolleri>", *(kk or ["- Kontrol bulgusu yok"]), "</kod_kontrolleri>"]), terimler)
    yanit = llm.json_iste(sistem, mesaj, SEMA)
    bilinen = sayilar(mesaj)
    metin = " ".join([yanit["urun_tanimi"], yanit["baski_nakis"], yanit["bakim_talimati_onerisi"]]
                     + [x["detay"] for x in yanit["yapim_detaylari"]]
                     + [x for a in ("dikis_talimatlari", "etiket_yerlesimi", "utu_paketleme", "kalite_notlari") for x in yanit[a]])
    # Ölçü/sıcaklık gibi birimli sayılar girdide yoksa işaretle (dikiş sıklığı gibi atölye standartları için de kontrol edilsin)
    yanit["dogrulanamayan_sayilar"] = sorted({m.group(0).strip() for m in re.finditer(
        r"(\d+(?:[.,]\d+)?)\s*(cm|mm|°c|derece|g/m²|gr|%|iğne|iplik|adım)", metin, re.I) if sayilar(m.group(1)) - bilinen})
    return yanit


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
ONEM_DOLGU = {"yüksek": PatternFill("solid", fgColor="FDE2E1"), "orta": PatternFill("solid", fgColor="FFF4CE"),
              "bilgi": PatternFill("solid", fgColor="E8F0FE")}
TASLAK_DOLGU = PatternFill("solid", fgColor="FFF8E1")
ONAY_DOLGU = PatternFill("solid", fgColor="FFF4CE")
UST = Alignment(vertical="top", wrap_text=True)
INCE = Side(style="thin", color="BBBBBB")
CERCEVE = Border(left=INCE, right=INCE, top=INCE, bottom=INCE)


def _baslik(ws, satir=1):
    for h in ws[satir]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI


def _hizala(ws):
    for r in ws.iter_rows():
        for h in r:
            h.alignment = UST


def rapor_yaz(cikti: Path, klasor: Path, metinler, atlanan, k: dict, y: dict, terimler: list[str]) -> Path:
    g = lambda s: geri_ac(s, terimler)  # noqa: E731
    mb = k["model"]
    baslik = f"{mb['kod'] or klasor.name} — {mb['ad'] or ''}".strip(" —")
    wb = Workbook()
    kp = wb.active
    kp.title = "Kapak"
    kp.append(["TEKNİK FÖY (TECH PACK) — TASLAK"])
    kp["A1"].font = Font(bold=True, size=14)
    for satir in [["Model kodu", mb["kod"]], ["Model adı", mb["ad"]], ["Sezon", mb["sezon"]], ["Müşteri / marka", mb["musteri"]],
                  ["Bedenler", ", ".join(k["olcu"]["bedenler"]) if k["olcu"] else ", ".join(mb["bedenler"])],
                  ["Ana (numune) beden", mb["ana_beden"]], ["Renkler", ", ".join(mb["renkler"])], ["Ürün tanımı", g(y["urun_tanimi"])],
                  ["Çizim / görseller", ", ".join(a for a in atlanan if re.search(r"\.(png|jpe?g|ai|svg|gif|webp)\b", a, re.I)) or "—"],
                  ["Hazırlanma", datetime.now().strftime("%d.%m.%Y")], ["Revizyon", "0 (taslak)"], []]:
        kp.append(satir)
    kp.append(["Onaylar", "Ad Soyad", "Tarih", "İmza"])
    _baslik(kp, kp.max_row)
    for rol in ("Tasarımcı", "Modelist", "Kalite / Ürün geliştirme", "Müşteri"):
        kp.append([rol, "", "", ""])
        for h in kp[kp.max_row]:
            h.border = CERCEVE
    kp.column_dimensions["A"].width = 24
    kp.column_dimensions["B"].width = 90
    kp.column_dimensions["C"].width = 14
    kp.column_dimensions["D"].width = 18
    _hizala(kp)

    if k["olcu"]:
        ol = wb.create_sheet("Ölçü Tablosu")
        bd = k["olcu"]["bedenler"]
        ol.append(["Kod", "Ölçü Noktası", "Tolerans ±"] + bd)
        _baslik(ol)
        for r in k["olcu"]["satirlar"]:
            ol.append([r["kod"], r["ad"], float(r["tol"]) if r["tol"] is not None else None]
                      + [float(r["deger"][x]) if r["deger"].get(x) is not None else None for x in bd])
            if mb["ana_beden"] in bd:
                ol.cell(ol.max_row, 4 + bd.index(mb["ana_beden"])).font = Font(bold=True)
        ol.append([])
        ol.append(["", "Ölçüler cm'dir; yarım ölçüler düz zeminde ölçülür. Ana beden kalın yazılmıştır."])
        ol.column_dimensions["A"].width = 6
        ol.column_dimensions["B"].width = 48
        for j in range(3, len(bd) + 4):
            ol.column_dimensions[get_column_letter(j)].width = 9
        ol.freeze_panes = "C2"

    if k["bom"]:
        bm = wb.create_sheet("Malzeme Listesi")
        bm.append(["Tür", "Malzeme Kodu", "Malzeme", "Kompozisyon", "Gramaj", "En", "Renk", "Tedarikçi", "Kullanım Yeri", "Tüketim", "Birim"])
        _baslik(bm)
        for x in k["bom"]:
            bm.append([x["tur"], x["kod"], x["ad"], x["komp"], x["gramaj"], x["en"], x["renk"], g(x["tedarikci"]), x["yer"], x["tuketim"], x["birim"]])
        for j, w in enumerate((11, 13, 32, 26, 9, 7, 14, 18, 26, 9, 7), 1):
            bm.column_dimensions[get_column_letter(j)].width = w
        bm.freeze_panes = "D2"

    yd = wb.create_sheet("Yapım Detayları")
    yd.append(["Bölüm", "Detay (taslak — modelist onayı gerekir)"])
    _baslik(yd)
    for x in y["yapim_detaylari"]:
        yd.append([g(x["bolum"]), g(x["detay"])])
        yd.cell(yd.max_row, 2).fill = TASLAK_DOLGU
    yd.column_dimensions["A"].width = 22
    yd.column_dimensions["B"].width = 110
    _hizala(yd)

    t = wb.create_sheet("Talimatlar")
    t.append(["Başlık", "Talimat (taslak)"])
    _baslik(t)
    for bas, liste in (("Dikiş", y["dikis_talimatlari"]), ("Baskı / nakış", [y["baski_nakis"]] if y["baski_nakis"] else []),
                       ("Etiket yerleşimi", y["etiket_yerlesimi"]), ("Ütü ve paketleme", y["utu_paketleme"]),
                       ("Kalite notları", y["kalite_notlari"]), ("Bakım talimatı önerisi", [y["bakim_talimati_onerisi"]]),
                       ("Açık sorular", y["acik_sorular"])):
        for i, x in enumerate(liste):
            t.append([bas if i == 0 else "", g(x)])
            t.cell(t.max_row, 2).fill = TASLAK_DOLGU
    t.append([])
    t.append(["Not", "Bakım talimatı ve sembolleri (ISO 3758) kumaş testleriyle doğrulanmadan etikete basılmamalıdır."])
    t.column_dimensions["A"].width = 22
    t.column_dimensions["B"].width = 110
    _hizala(t)

    ko = wb.create_sheet("Kontroller")
    ko.append(["Önem", "Bulgu", "Kaynak", "Üreten", "Onay / Düzeltme"])
    _baslik(ko)
    for o, a, kay in k["bulgular"]:
        ko.append([o, g(a), kay, "Kod", ""])
        ko.cell(ko.max_row, 1).fill = ONEM_DOLGU[o]
        ko.cell(ko.max_row, 5).fill = ONAY_DOLGU
    for s in y["dogrulanamayan_sayilar"]:
        ko.append(["orta", f"Model metninde girdilerde olmayan değer: '{s}' (atölye standardı mı, müşteri talebi mi? doğrulayın)", "model metni", "Kod", ""])
        ko.cell(ko.max_row, 1).fill = ONEM_DOLGU["orta"]
        ko.cell(ko.max_row, 5).fill = ONAY_DOLGU
    for j, w in enumerate((9, 100, 26, 8, 22), 1):
        ko.column_dimensions[get_column_letter(j)].width = w
    _hizala(ko)

    rv = wb.create_sheet("Revizyonlar")
    rv.append(["Rev", "Tarih", "Sayfa / Bölüm", "Değişiklik", "Yapan", "Onaylayan"])
    _baslik(rv)
    rv.append([0, datetime.now().strftime("%d.%m.%Y"), "Tümü", "İlk taslak (Workers / Workless agent)", "", ""])
    for j, w in enumerate((5, 12, 18, 60, 16, 16), 1):
        rv.column_dimensions[get_column_letter(j)].width = w
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

    md = [f"# Teknik föy taslağı — {baslik}", "",
          f"**Sezon:** {mb['sezon'] or '—'} · **Bedenler:** {', '.join(k['olcu']['bedenler']) if k['olcu'] else '—'}"
          f" · **Renkler:** {', '.join(mb['renkler']) or '—'}", "", "## Ürün tanımı", g(y["urun_tanimi"]), "", "## Yapım detayları"]
    md += [f"- **{g(x['bolum'])}:** {g(x['detay'])}" for x in y["yapim_detaylari"]]
    md += ["", "## Dikiş talimatları", *[f"- {g(x)}" for x in y["dikis_talimatlari"]]]
    if y["baski_nakis"]:
        md += ["", "## Baskı / nakış", g(y["baski_nakis"])]
    md += ["", "## Etiket yerleşimi", *[f"- {g(x)}" for x in y["etiket_yerlesimi"]], "", "## Ütü ve paketleme",
           *[f"- {g(x)}" for x in y["utu_paketleme"]], "", "## Kalite notları", *[f"- {g(x)}" for x in y["kalite_notlari"]], "",
           "## Bakım talimatı önerisi (test edilmeli)", g(y["bakim_talimati_onerisi"]), "", "## Kontrol bulguları"]
    md += [f"- **[{o}]** {g(a)} — _kod kontrolü_" for o, a, _ in k["bulgular"]] or ["- Bulgu yok"]
    if y["dogrulanamayan_sayilar"]:
        md += [f"- **[orta]** Model metninde girdilerde olmayan değerler: {', '.join(y['dogrulanamayan_sayilar'])} — doğrulayın"]
    md += ["", "## Açık sorular", *[f"- {g(x)}" for x in y["acik_sorular"]], "", "---",
           f"_Ölçü tablosu ve malzeme listesi Excel föyde: {cikti.name}. Kaynaklar: {', '.join(d.ad for d in metinler)}_",
           "_Yapay zekâ destekli taslak. Ölçüler ve malzemeler girdilerden aynen aktarılmıştır; yapım detayları ve talimatlar "
           "modelist ve kalite onayıyla kesinleşir._"]
    md_yol = cikti.with_suffix(".md")
    md_yol.write_text("\n".join(md), encoding="utf-8")
    return md_yol


def calistir(klasor: Path, cikti: Path, terimler: list[str] | None = None, evet: bool = False) -> dict:
    if not klasor.is_dir():
        raise llm.LLMHatasi(f"{klasor}: klasör bulunamadı (model dosyalarını bir klasörde toplayın).")
    g, metinler, atlanan = dosyalari_oku(klasor)
    if not metinler:
        raise llm.LLMHatasi(f"{klasor}: okunabilir dosya yok.")
    k = kontrol_et(g, metinler)
    print(f"[OK] {len(metinler)} dosya · {len(k['olcu']['satirlar']) if k['olcu'] else 0} ölçü noktası · {len(k['bom'])} malzeme · "
          f"kod kontrolü: {len(k['bulgular'])} bulgu")
    for o, a, _ in k["bulgular"]:
        print(f"[!] [{o}] {a}")
    terimler = terimler or []
    llm.onay_al(f"Model bilgisi, ölçü tablosu ve malzeme listesi ({len(terimler)} gizli terim maskeli) teknik föy taslağı için gönderilecek.", evet)
    y = model_yaz(metinler, k, terimler)
    md = rapor_yaz(cikti, klasor, metinler, atlanan, k, y, terimler)
    return {"girdi": g, "kontrol": k, "yanit": y, "md": md, "atlanan": atlanan}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    ornek = BURASI / "ornek_veri" / "TS-2027-014"
    p = argparse.ArgumentParser(description="Model bilgisi, ölçü tablosu ve malzeme listesinden teknik föy (tech pack) taslağı hazırlar.")
    p.add_argument("--girdi", type=Path, default=ornek, help="Model klasörü (model bilgisi .txt, ölçü tablosu, malzeme listesi)")
    p.add_argument("--gizle", nargs="*", default=[], metavar="TERİM", help="Modele gönderilmeden maskelenecek adlar (müşteri markası, tedarikçi)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "teknik_foy.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    if a.girdi == ornek:
        a.gizle = a.gizle or ["Kurgu Moda"]
    try:
        s = calistir(a.girdi, a.cikti, a.gizle, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    if s["yanit"]["dogrulanamayan_sayilar"]:
        print(f"[!] Model metninde girdilerde olmayan değerler: {', '.join(s['yanit']['dogrulanamayan_sayilar'])}")
    print(f"[OK] Teknik föy: {a.cikti.resolve()}")
    print(f"[OK] Özet: {s['md'].resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
