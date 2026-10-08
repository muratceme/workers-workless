"""
Performans Değerlendirme Özeti — Workers / Workless AI Agent
İnsan Kaynakları › İnsan Kaynakları Müdürü

1. Kod, değerlendirme formlarını (hedef ve yetkinlik kalemleri, ağırlık, yönetici puanı, öz değerlendirme,
   yorum) okur ve tüm puanları kendisi hesaplar: hedef puanı, yetkinlik puanı, ağırlıklı genel puan ve
   performans kategorisi; departman ve yönetici bazında dağılım; değerlendirici eğilimi (cömert/katı,
   merkeze yığılma); öz değerlendirme ile yönetici puanı arasındaki büyük farklar; ağırlık toplamı hataları.
2. Yorumlar ortak ayrımcılık tarayıcısından geçer; korunan özelliğe (yaş, cinsiyet, gebelik, medeni hâl, sağlık,
   din, köken…) dayalı yorumlar işaretlenir ve modele "DEĞERLENDİRME DIŞI" etiketiyle gider.
3. Çalışan ve yönetici adları modele takma adla (P1, Y1) gider. Model, her çalışan için yorumlara dayanarak
   güçlü yönler, gelişim alanları, gelişim önerileri ve geri bildirim görüşmesinde konuşulacakları yazar.
   Puan ve kategori modelden değil koddan gelir.
4. Sonuç Excel'e yazılır; her özet için 'Onay' sütunu ve yazdırılabilir özet kartları vardır.

Kullanım:
    python agent.py                                              # örnek formlarla dener
    python agent.py --girdi degerlendirmeler.xlsx --hedef-agirligi 60
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import Counter, OrderedDict, defaultdict
from datetime import datetime
from pathlib import Path
from statistics import fmean, pstdev

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import ayrimcilik
import llm

BURASI = Path(__file__).resolve().parent
KATEGORILER = [  # (alt sınır, ad) — 1-5 ölçeği
    (4.5, "Üstün"), (3.8, "Beklentinin üzerinde"), (3.0, "Beklentiyi karşılıyor"), (2.0, "Gelişime açık"), (0.0, "Beklentinin altında")]
EGILIM_ESIK = 0.5          # yönetici ortalaması ile şirket ortalaması farkı
YIGILMA_ESIK = 0.25        # yöneticinin verdiği genel puanların standart sapması (≥ 4 çalışanda)
OZ_FARK = 1.5              # öz değerlendirme ile yönetici puanı farkı (kalem bazında)


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def katla(s) -> str:
    s = kucuk(s)
    for a, b in zip("ıöüşğç", "iousgc"):
        s = s.replace(a, b)
    return re.sub(r"[^a-z0-9%]+", " ", s).strip()


def puan(x) -> float | None:
    if x in (None, ""):
        return None
    try:
        return float(str(x).replace(",", "."))
    except ValueError:
        return None


def yaz(x: float | None) -> str:
    return "-" if x is None else f"{x:.2f}".replace(".", ",")


def kategori(p: float | None) -> str:
    if p is None:
        return "-"
    return next(ad for alt, ad in KATEGORILER if p >= alt)


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:5])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


# ----------------------------------------------------------------------------
# Okuma ve hesap
# ----------------------------------------------------------------------------

def formlari_oku(yol: Path) -> list[dict]:
    s = tablo_oku(yol)
    kb = [katla(x) for x in s[0]]
    bul = lambda *a: next((kb.index(katla(x)) for x in a if katla(x) in kb), None)  # noqa: E731
    i = {"sicil": bul("sicil", "sicil no"), "ad": bul("ad soyad", "çalışan", "personel"), "dep": bul("departman", "birim"),
         "yon": bul("yönetici", "değerlendiren", "amir"), "tur": bul("tür", "kalem türü", "bölüm"),
         "kalem": bul("kalem", "hedef yetkinlik", "hedef", "yetkinlik", "kriter"), "agirlik": bul("ağırlık", "ağırlık %"),
         "puan": bul("puan", "yönetici puanı", "değerlendirme puanı"), "oz": bul("öz puan", "öz değerlendirme", "öz değerlendirme puanı"),
         "yorum": bul("yorum", "açıklama", "yönetici yorumu", "not")}
    for k, ad in (("ad", "Ad Soyad"), ("kalem", "Kalem"), ("puan", "Puan")):
        if i[k] is None:
            raise llm.LLMHatasi(f"Değerlendirme dosyasında '{ad}' sütunu bulunamadı. Başlıklar: {s[0]}")
    al = lambda r, k: r[i[k]] if i[k] is not None and i[k] < len(r) else None  # noqa: E731
    satirlar = []
    for r in s[1:]:
        ad = str(al(r, "ad") or "").strip()
        if not ad:
            continue
        tur = katla(al(r, "tur"))
        yorum = str(al(r, "yorum") or "").strip()
        satirlar.append({"sicil": str(al(r, "sicil") or ad).strip(), "ad": ad, "dep": str(al(r, "dep") or "-").strip(),
                         "yon": str(al(r, "yon") or "-").strip(), "tur": "Hedef" if tur.startswith(("hedef", "okr", "kpi")) else "Yetkinlik",
                         "kalem": str(al(r, "kalem") or "").strip(), "agirlik": puan(al(r, "agirlik")), "puan": puan(al(r, "puan")),
                         "oz": puan(al(r, "oz")), "yorum": yorum,
                         "bulgular": [b for b in ayrimcilik.tara(yorum) if b["seviye"] == "yüksek"] if yorum else []})
    return satirlar


def agirlikli(kalemler: list[dict]) -> tuple[float | None, list[str]]:
    puanli = [k for k in kalemler if k["puan"] is not None]
    if not puanli:
        return None, []
    uyari = []
    if all(k["agirlik"] is not None for k in puanli):
        toplam = sum(k["agirlik"] for k in puanli)
        if abs(toplam - 100) > 0.5 and abs(toplam - 1) > 0.005:
            uyari.append(f"ağırlık toplamı {toplam:g} (100 olmalı); oransal ölçeklendi")
        return sum(k["puan"] * k["agirlik"] for k in puanli) / toplam if toplam else None, uyari
    return fmean(k["puan"] for k in puanli), uyari


def hesapla(satirlar: list[dict], hedef_agirligi: float) -> dict:
    calisanlar: OrderedDict[str, dict] = OrderedDict()
    for r in satirlar:
        c = calisanlar.setdefault(r["sicil"], {"sicil": r["sicil"], "ad": r["ad"], "dep": r["dep"], "yon": r["yon"], "kalemler": []})
        c["kalemler"].append(r)
    for c in calisanlar.values():
        h = [k for k in c["kalemler"] if k["tur"] == "Hedef"]
        y = [k for k in c["kalemler"] if k["tur"] == "Yetkinlik"]
        c["hedef"], u1 = agirlikli(h)
        c["yetkinlik"], u2 = agirlikli(y)
        c["uyarilar"] = [f"Hedef: {u}" for u in u1] + [f"Yetkinlik: {u}" for u in u2]
        if c["hedef"] is not None and c["yetkinlik"] is not None:
            c["genel"] = c["hedef"] * hedef_agirligi / 100 + c["yetkinlik"] * (1 - hedef_agirligi / 100)
        else:
            c["genel"] = c["hedef"] if c["hedef"] is not None else c["yetkinlik"]
        c["kategori"] = kategori(c["genel"])
        eksik = [k["kalem"] for k in c["kalemler"] if k["puan"] is None]
        if eksik:
            c["uyarilar"].append("puanlanmamış kalem: " + ", ".join(eksik))
        c["oz_farklar"] = [k for k in c["kalemler"] if k["puan"] is not None and k["oz"] is not None and abs(k["puan"] - k["oz"]) >= OZ_FARK]
        c["onyargi"] = [(k["kalem"], b) for k in c["kalemler"] for b in k["bulgular"]]
    genel = [c["genel"] for c in calisanlar.values() if c["genel"] is not None]
    sirket_ort = fmean(genel) if genel else None
    yoneticiler = defaultdict(list)
    for c in calisanlar.values():
        if c["genel"] is not None:
            yoneticiler[c["yon"]].append(c["genel"])
    egilim = {}
    for y, ps in yoneticiler.items():
        ort = fmean(ps)
        notlar = []
        if sirket_ort is not None and len(ps) >= 2 and ort - sirket_ort >= EGILIM_ESIK:
            notlar.append("cömert değerlendirme eğilimi olabilir")
        if sirket_ort is not None and len(ps) >= 2 and sirket_ort - ort >= EGILIM_ESIK:
            notlar.append("katı değerlendirme eğilimi olabilir")
        if len(ps) >= 4 and pstdev(ps) < YIGILMA_ESIK:
            notlar.append("puanlar ortada yığılmış (ayrıştırma zayıf)")
        egilim[y] = {"n": len(ps), "ort": ort, "sd": pstdev(ps) if len(ps) > 1 else 0.0, "notlar": notlar}
    return {"calisanlar": calisanlar, "sirket_ort": sirket_ort, "yoneticiler": egilim}


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

SEMA = {
    "type": "object",
    "properties": {
        "calisanlar": {"type": "array", "items": {"type": "object", "properties": {
            "calisan": {"type": "string"},
            "guclu_yonler": {"type": "array", "items": {"type": "string"}},
            "gelisim_alanlari": {"type": "array", "items": {"type": "string"}},
            "gelisim_onerileri": {"type": "array", "items": {"type": "string"}},
            "gorusme_notlari": {"type": "array", "items": {"type": "string"}},
            "ozet": {"type": "string"}},
            "required": ["calisan", "guclu_yonler", "gelisim_alanlari", "gelisim_onerileri", "gorusme_notlari", "ozet"],
            "additionalProperties": False}}},
    "required": ["calisanlar"],
    "additionalProperties": False,
}


def takma(h: dict) -> dict[str, str]:
    harita = {}
    for i, c in enumerate(h["calisanlar"].values(), 1):
        harita[c["ad"]] = f"P{i}"
    for i, y in enumerate(sorted({c["yon"] for c in h["calisanlar"].values() if c["yon"] != "-"}), 1):
        harita[y] = f"Y{i}"
    return harita


def gizle(metin: str, harita: dict[str, str]) -> str:
    metin = llm.maskele(metin)
    for gercek in sorted(harita, key=len, reverse=True):
        metin = re.sub(re.escape(gercek), harita[gercek], metin, flags=re.I)
        ilk = gercek.split()[0]
        if len(ilk) >= 3:                                   # yalnız adla anılmışsa da gizle
            metin = re.sub(rf"\b{re.escape(ilk)}\b", harita[gercek], metin, flags=re.I)
    return metin


def ac(metin: str, harita: dict[str, str]) -> str:
    for gercek, t in sorted(harita.items(), key=lambda i: -len(i[1])):
        metin = re.sub(rf"\b{re.escape(t)}\b", gercek, metin)
    return metin


def ozetle(h: dict, harita: dict[str, str]) -> dict[str, dict]:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    sonuc = {}
    calisanlar = list(h["calisanlar"].values())
    for i in range(0, len(calisanlar), 10):
        satirlar = []
        for c in calisanlar[i:i + 10]:
            satirlar.append(f'<calisan id="{harita[c["ad"]]}" genel="{yaz(c["genel"])}" kategori="{c["kategori"]}" '
                            f'hedef="{yaz(c["hedef"])}" yetkinlik="{yaz(c["yetkinlik"])}">')
            for k in c["kalemler"]:
                oz = f" | öz {k['oz']:g}" if k["oz"] is not None else ""
                etiket = "[DEĞERLENDİRME DIŞI] " if k["bulgular"] else ""
                satirlar.append(f"{k['tur']} | {k['kalem']} | puan {k['puan'] if k['puan'] is not None else '-'}{oz} | "
                                f"{etiket}{gizle(k['yorum'], harita) or '-'}")
            satirlar.append("</calisan>")
        yanit = llm.json_iste(sistem, "<calisanlar>\n" + "\n".join(satirlar) + "\n</calisanlar>", SEMA)
        geri = {v: k for k, v in harita.items()}
        for s in yanit.get("calisanlar", []):
            if s.get("calisan") in geri:
                sonuc[geri[s["calisan"]]] = s
    return sonuc


def denetle(ozet: dict) -> list[str]:
    """Model metninde korunan özelliğe dayalı ifade kalmışsa işaretle."""
    metin = " ".join([ozet["ozet"], *ozet["guclu_yonler"], *ozet["gelisim_alanlari"], *ozet["gelisim_onerileri"]])
    return [f"özet metninde riskli ifade: '{b['ifade']}' ({b['kategori']})" for b in ayrimcilik.tara(metin) if b["seviye"] == "yüksek"]


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
ONAY_DOLGU = PatternFill("solid", fgColor="FFF4CE")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
UST = Alignment(vertical="top", wrap_text=True)
KAT_DOLGU = {"Üstün": "C9E7C4", "Beklentinin üzerinde": "E3F5E1", "Beklentiyi karşılıyor": "FFFFFF", "Gelişime açık": "FFF4CE",
             "Beklentinin altında": "FDE2E1", "-": "FFFFFF"}


def _baslik(ws, satir=1):
    for c in ws[satir]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI


def madde(liste: list[str], harita: dict) -> str:
    return "\n".join(f"• {ac(x, harita)}" for x in liste)


def rapor_yaz(cikti: Path, h: dict, ozet: dict, kontrol: dict, harita: dict, hedef_agirligi: float) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Çalışan Özeti"
    ws.append(["Sicil", "Ad Soyad", "Departman", "Yönetici", "Hedef Puanı", "Yetkinlik Puanı", "Genel Puan", "Kategori",
               "Güçlü Yönler", "Gelişim Alanları", "Gelişim Önerileri", "Görüşmede Konuşulacaklar", "Özet", "Uyarılar", "İK Onayı"])
    _baslik(ws)
    for c in sorted(h["calisanlar"].values(), key=lambda c: -(c["genel"] or 0)):
        o = ozet.get(c["sicil"]) or ozet.get(c["ad"]) or {}
        uyarilar = list(c["uyarilar"])
        uyarilar += [f"değerlendirme dışı yorum ({kalem}): '{b['ifade']}' — {b['kategori']}" for kalem, b in c["onyargi"]]
        uyarilar += [f"öz değerlendirme farkı: {k['kalem']} (yönetici {k['puan']:g}, öz {k['oz']:g})" for k in c["oz_farklar"]]
        uyarilar += kontrol.get(c["ad"], [])
        ws.append([c["sicil"], c["ad"], c["dep"], c["yon"], c["hedef"], c["yetkinlik"], c["genel"], c["kategori"],
                   madde(o.get("guclu_yonler", []), harita), madde(o.get("gelisim_alanlari", []), harita),
                   madde(o.get("gelisim_onerileri", []), harita), madde(o.get("gorusme_notlari", []), harita),
                   ac(o.get("ozet", "AI özet üretmedi"), harita), "\n".join(uyarilar), ""])
        n = ws.max_row
        for col in (5, 6, 7):
            ws.cell(n, col).number_format = "0.00"
        ws.cell(n, 8).fill = PatternFill("solid", fgColor=KAT_DOLGU[c["kategori"]])
        for col in (9, 10, 11, 12, 13):
            ws.cell(n, col).fill = AI_DOLGU
        if any("değerlendirme dışı" in u or "riskli ifade" in u for u in uyarilar):
            ws.cell(n, 14).fill = KIRMIZI
        ws.cell(n, 15).fill = ONAY_DOLGU
        for x in ws[n]:
            x.alignment = UST
    for j, w in enumerate((9, 20, 14, 16, 9, 9, 9, 18, 40, 40, 40, 40, 50, 45, 10), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    d = wb.create_sheet("Dağılım")
    kat_adlari = [ad for _, ad in reversed(KATEGORILER)]
    d.append(["Kategori", "Çalışan", "Pay"])
    _baslik(d)
    toplam = len(h["calisanlar"])
    say = Counter(c["kategori"] for c in h["calisanlar"].values())
    for k in kat_adlari:
        d.append([k, say.get(k, 0), say.get(k, 0) / toplam if toplam else None])
        d.cell(d.max_row, 3).number_format = "0%"
        d.cell(d.max_row, 1).fill = PatternFill("solid", fgColor=KAT_DOLGU[k])
    g = BarChart()
    g.title, g.height, g.width = "Performans kategorisi dağılımı", 7, 16
    g.add_data(Reference(d, min_col=2, min_row=1, max_row=1 + len(kat_adlari)), titles_from_data=True)
    g.set_categories(Reference(d, min_col=1, min_row=2, max_row=1 + len(kat_adlari)))
    d.add_chart(g, "E2")
    d.append([])
    d.append(["Departman"] + kat_adlari + ["Ortalama"])
    _baslik(d, d.max_row)
    deps = defaultdict(list)
    for c in h["calisanlar"].values():
        deps[c["dep"]].append(c)
    for dep, cs in sorted(deps.items()):
        s = Counter(c["kategori"] for c in cs)
        ps = [c["genel"] for c in cs if c["genel"] is not None]
        d.append([dep] + [s.get(k, 0) or None for k in kat_adlari] + [round(fmean(ps), 2) if ps else None])
    d.column_dimensions["A"].width = 24
    for j in range(2, 8):
        d.column_dimensions[get_column_letter(j)].width = 14

    y = wb.create_sheet("Değerlendiriciler")
    y.append(["Yönetici", "Değerlendirilen", "Ortalama Genel Puan", "Şirket Ortalaması", "Fark", "Std. Sapma", "Not"])
    _baslik(y)
    for ad, v in sorted(h["yoneticiler"].items()):
        y.append([ad, v["n"], round(v["ort"], 2), round(h["sirket_ort"], 2) if h["sirket_ort"] is not None else None,
                  round(v["ort"] - h["sirket_ort"], 2) if h["sirket_ort"] is not None else None, round(v["sd"], 2), "; ".join(v["notlar"])])
        if v["notlar"]:
            y.cell(y.max_row, 7).fill = ONAY_DOLGU
    y.append([])
    y.append(["Kalibrasyon toplantısında değerlendiriciler arası tutarlılığı tartışın; bu tablo bir yöneticinin hatalı "
              "olduğunu göstermez, yalnızca istatistiksel bir sinyaldir (ekipler gerçekten farklı performans gösterebilir)."])
    for j, w in enumerate((20, 14, 18, 16, 8, 10, 45), 1):
        y.column_dimensions[get_column_letter(j)].width = w

    k = wb.create_sheet("Kalem Detayı")
    k.append(["Sicil", "Ad Soyad", "Tür", "Kalem", "Ağırlık", "Yönetici Puanı", "Öz Puan", "Yorum", "Değerlendirme Dışı"])
    _baslik(k)
    for c in h["calisanlar"].values():
        for x in c["kalemler"]:
            k.append([c["sicil"], c["ad"], x["tur"], x["kalem"], x["agirlik"], x["puan"], x["oz"], x["yorum"],
                      ayrimcilik.ozet(x["bulgular"]) if x["bulgular"] else ""])
            if x["bulgular"]:
                k.cell(k.max_row, 9).fill = KIRMIZI
            if x["oz"] is not None and x["puan"] is not None and abs(x["puan"] - x["oz"]) >= OZ_FARK:
                k.cell(k.max_row, 7).fill = ONAY_DOLGU
    for j, w in enumerate((9, 20, 10, 28, 8, 9, 8, 60, 40), 1):
        k.column_dimensions[get_column_letter(j)].width = w
    k.freeze_panes = "C2"
    k.auto_filter.ref = k.dimensions

    kr = wb.create_sheet("Özet Kartları")
    kr.column_dimensions["A"].width = 26
    kr.column_dimensions["B"].width = 100
    for c in sorted(h["calisanlar"].values(), key=lambda c: (c["dep"], c["ad"])):
        o = ozet.get(c["sicil"]) or ozet.get(c["ad"]) or {}
        kr.append([f"{c['ad']} · {c['dep']}", f"Genel {yaz(c['genel'])} — {c['kategori']} (hedef {yaz(c['hedef'])}, yetkinlik {yaz(c['yetkinlik'])})"])
        for x in kr[kr.max_row]:
            x.fill, x.font = BASLIK_DOLGU, BASLIK_YAZI
        for baslik, alan in (("Güçlü yönler", "guclu_yonler"), ("Gelişim alanları", "gelisim_alanlari"),
                             ("Gelişim önerileri", "gelisim_onerileri"), ("Görüşmede konuşulacaklar", "gorusme_notlari")):
            kr.append([baslik, madde(o.get(alan, []), harita)])
            kr.cell(kr.max_row, 2).alignment = UST
        kr.append([])

    b = wb.create_sheet("Bilgi")
    for s in [["Genel puan", f"hedef puanı × %{hedef_agirligi:g} + yetkinlik puanı × %{100 - hedef_agirligi:g}; kalem puanları ağırlıklı ortalama (1–5)"],
              ["Kategoriler", " · ".join(f"{ad} ≥ {alt:g}".replace(".", ",") for alt, ad in KATEGORILER if alt > 0) + " · Beklentinin altında < 2"],
              ["Değerlendirici", f"yönetici ortalaması şirket ortalamasından ±{EGILIM_ESIK:g} farklıysa eğilim; 4+ çalışanda std. sapma < "
                                 f"{YIGILMA_ESIK:g} ise ortada yığılma".replace(".", ",")],
              ["Öz değerlendirme", f"aynı kalemde yönetici ile öz puan farkı ≥ {OZ_FARK:g} ise görüşmede ele alınmalı".replace(".", ",")],
              ["Değerlendirme dışı", "korunan özelliğe (yaş, cinsiyet, gebelik, medeni hâl, sağlık, din, köken…) dayalı yorumlar "
                                     "(İş K. md. 5, 6701 s. Kanun) işaretlenir; puanı yeniden gözden geçirin"],
              ["Gizlilik", "Çalışan ve yönetici adları modele takma adla (P1, Y1) gönderildi; rapor şirket içi olduğundan gerçek adlar gösterilir"],
              ["Model", llm.kullanim_ozeti()], ["Oluşturulma", datetime.now().strftime("%d.%m.%Y %H:%M")],
              ["Önemli", "Puan ve kategoriler koddan, metinler modelden gelir. Özetler taslaktır; İK ve yönetici onayı olmadan çalışana "
                         "iletilmemeli ve ücret/terfi kararlarında tek başına kullanılmamalıdır."]]:
        b.append(s)
    b.column_dimensions["A"].width = 20
    b.column_dimensions["B"].width = 130
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(girdi: Path, cikti: Path, hedef_agirligi: float = 60.0, evet: bool = False) -> dict:
    satirlar = formlari_oku(girdi)
    if not satirlar:
        raise llm.LLMHatasi(f"{girdi}: değerlendirme satırı yok.")
    h = hesapla(satirlar, hedef_agirligi)
    harita = takma(h)
    onyargi = sum(len(c["onyargi"]) for c in h["calisanlar"].values())
    print(f"[OK] {len(h['calisanlar'])} çalışan · {len(h['yoneticiler'])} yönetici · şirket ortalaması {yaz(h['sirket_ort'])}"
          + (f" · değerlendirme dışı yorum: {onyargi}" if onyargi else ""))
    llm.onay_al("Değerlendirme puanları ve yorumları (çalışan/yönetici adları takma adlı; e-posta/telefon maskeli) özet için gönderilecek.", evet)
    ham = ozetle(h, harita)
    ozet = {}
    kontrol = {}
    for ad, o in ham.items():
        sicil = next(c["sicil"] for c in h["calisanlar"].values() if c["ad"] == ad)
        ozet[sicil] = o
        kontrol[ad] = denetle(o)
    rapor_yaz(cikti, h, ozet, kontrol, harita, hedef_agirligi)
    return {"h": h, "ozet": ozet, "kontrol": kontrol, "harita": harita}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Performans değerlendirme formlarını çalışan başına özete ve dağılım raporuna çevirir.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "degerlendirmeler.csv",
                   help="Formlar (.xlsx/.csv): Sicil, Ad Soyad, Departman, Yönetici, Tür (Hedef/Yetkinlik), Kalem, Ağırlık, Puan [, Öz Puan], Yorum")
    p.add_argument("--hedef-agirligi", type=float, default=60.0, help="Genel puanda hedeflerin ağırlığı, %% (varsayılan 60)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "performans_ozeti.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.cikti, a.hedef_agirligi, a.evet)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(s['ozet'])}/{len(s['h']['calisanlar'])} çalışan özetlendi")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
