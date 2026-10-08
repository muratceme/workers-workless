"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri" / "degerlendirmeler.csv"
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    sonuc = []
    for i in re.findall(r'<calisan id="(P\d+)"', kullanici):
        ozet = f"{i} hedeflerini büyük ölçüde gerçekleştirdi."
        if i == "P3":
            ozet = "Hamile olduğu için müşteri ziyaretleri azaldı."       # model hatası: korunan özellik → kod yakalamalı
        sonuc.append({"calisan": i, "guclu_yonler": [f"{i} müşteri ilişkileri güçlü"], "gelisim_alanlari": ["planlama"],
                      "gelisim_onerileri": ["haftalık plan şablonu"], "gorusme_notlari": ["beklentileri netleştirin"], "ozet": ozet})
    return {"calisanlar": sonuc}


class HesapTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.h = agent.hesapla(agent.formlari_oku(ORNEK), 60)
        cls.c = cls.h["calisanlar"]

    def test_agirlikli_puan(self):
        e = self.c["1001"]
        self.assertAlmostEqual(e["hedef"], 5 * .5 + 4 * .3 + 4 * .2)                 # 4,5
        self.assertAlmostEqual(e["yetkinlik"], 5 * .4 + 4 * .3 + 4 * .3)             # 4,4
        self.assertAlmostEqual(e["genel"], 4.5 * .6 + 4.4 * .4)
        self.assertEqual(e["kategori"], "Beklentinin üzerinde")

    def test_agirlik_toplami_uyarisi(self):
        self.assertTrue(any("ağırlık toplamı 110" in u for u in self.c["1004"]["uyarilar"]))

    def test_onyargi_ve_oz_fark(self):
        self.assertTrue(self.c["1003"]["onyargi"])                                   # hamile / doğum izni
        self.assertTrue(self.c["1004"]["onyargi"])                                   # yaşı gereği
        self.assertFalse(self.c["1001"]["onyargi"])
        self.assertEqual([k["kalem"] for k in self.c["1002"]["oz_farklar"]], ["Tahsilat süresi (ort. 45 gün)", "Planlama"])

    def test_degerlendirici_egilimi(self):
        self.assertIn("cömert değerlendirme eğilimi olabilir", self.h["yoneticiler"]["Okan Er"]["notlar"])


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_uctan_uca(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(ORNEK, Path(tmp) / "p.xlsx", evet=True)
            wb = load_workbook(Path(tmp) / "p.xlsx")
            self.assertEqual(wb.sheetnames, ["Çalışan Özeti", "Dağılım", "Değerlendiriciler", "Kalem Detayı", "Özet Kartları", "Bilgi"])
            ozet_sayfa = {r[1]: r for r in wb["Çalışan Özeti"].iter_rows(min_row=2, values_only=True)}
            wb.close()
        gonderilen = "".join(MESAJLAR)
        for ad in ("Elif", "Arslan", "Murat Kılıç", "Zeynep"):
            self.assertNotIn(ad, gonderilen)                                         # adlar takma adla gider
        self.assertIn("[DEĞERLENDİRME DIŞI]", gonderilen)
        self.assertTrue(s["kontrol"]["Zeynep Kaya"])                                 # model metnindeki riskli ifade yakalandı
        self.assertIn("Elif Arslan müşteri ilişkileri güçlü", ozet_sayfa["Elif Arslan"][8])   # takma ad raporda geri açılır
        self.assertIn("riskli ifade", ozet_sayfa["Zeynep Kaya"][13])


if __name__ == "__main__":
    unittest.main()
