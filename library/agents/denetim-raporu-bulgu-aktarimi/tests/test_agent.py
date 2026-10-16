"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import json
import os
import sys
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

RAPOR = KLASOR / "ornek_veri" / "sosyal_denetim_raporu.txt"
METIN = RAPOR.read_text(encoding="utf-8")
CEVAP = json.loads((Path(__file__).parent / "sahte_cevaplar.json").read_text(encoding="utf-8"))
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    """İlk turda genç işçi bulgusunun derecesini rapordaki 'Majör'den 'Kritik'e yükseltir."""
    MESAJLAR.append(kullanici)
    return copy.deepcopy(CEVAP["dogru" if "<duzeltme>" in kullanici else "hatali"])


class KodTesti(unittest.TestCase):
    def test_termin_ve_esdeger(self):
        self.assertEqual(agent.termin_oku("kritik=3, minor=90"), {"Kritik": 3, "Majör": 30, "Minör": 90, "Belirtilmemiş": 30})
        with self.assertRaises(ValueError):
            agent.termin_oku("acil=1")
        self.assertEqual([agent.ESDEGER[agent.katla(x)] for x in ("Major", "Majör", "Critical", "Observation")], ["Majör", "Majör", "Kritik", "Gözlem"])
        self.assertTrue(all(agent.SIFIR.search(agent.katla(t)) for t in ("17 yaşında 2 çalışan", "Kimlik kartları işveren tarafından alıkonuldu",
                                                                       "Ekim ücretleri ödenmedi", "Acil çıkış kapısı kilitli kapı")))
        self.assertFalse(agent.SIFIR.search(agent.katla("Fazla mesai haftalık 66 saate ulaştı")))

    def test_denetim(self):
        self.assertEqual(agent.denetle(CEVAP["dogru"], METIN), [])
        self.assertEqual(agent.denetle(CEVAP["hatali"], METIN),
                         ["Bulgu 4: derece 'Kritik' rapordaki dereceyle ('Majör') aynı değil; raporun derecesini değiştirmeyin"])
        c = copy.deepcopy(CEVAP["dogru"])
        c["denetim_tarihi"] = "09.10.2026"
        c["bulgular"][0].update(alinti="dikim bölümünde 40 çalışan", bulgu="40 çalışan 70 saat çalıştı")
        c["bulgular"][1].update(rapordaki_derece="Critical")
        c["bulgular"][2].update(duzeltici_faaliyet="")
        c["bulgular"][4].update(rapordaki_derece="", derece="Gözlem")
        s = agent.denetle(c, METIN)
        for parca in ("denetim tarihi '09.10.2026'", "Bulgu 1: alıntı", "olmayan sayı: 40, 70", "'Critical' raporda geçmiyor", "Bulgu 3: düzeltici",
                      "'Belirtilmemiş' seçin"):
            self.assertTrue(any(parca in x for x in s), parca)

    def test_degerlendirme(self):
        b = agent.degerlendir(CEVAP["dogru"], agent.VARSAYILAN_TERMIN)
        self.assertEqual([(x["derece"], x["termin"], x["sifir"], x["takip"]) for x in b],
                         [("Majör", date(2026, 11, 7), False, True), ("Minör", date(2026, 12, 7), False, True), ("Majör", date(2026, 11, 7), False, True),
                          ("Majör", date(2026, 10, 15), True, True), ("Gözlem", None, False, False)])
        self.assertIn("Sıfır tolerans konusu (Genç işçiler)", b[3]["uyari"][0])


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ornek_ve_duzeltme(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "d.xlsx"
            self.assertEqual(agent.main(["--evet", "--onek", "SOS-26", "--cikti", str(cikti)]), 0)
            self.assertEqual(len(MESAJLAR), 2)
            wb = load_workbook(cikti)
            ws = wb["Bulgular"]
            self.assertEqual([c.value for c in ws[1]], agent.TAKIP)
            self.assertEqual(ws.max_row, 5)                                 # gözlem takip listesine girmez
            self.assertEqual([ws.cell(5, j).value for j in (1, 7, 10, 11)], ["SOS-26-04", "Majör", datetime(2026, 10, 15), "Açık"])
            self.assertEqual(ws.cell(2, 3).value, datetime(2026, 10, 8))
            d = [list(r) for r in wb["Ayrıntı"].iter_rows(min_row=2, values_only=True)]
            self.assertEqual([r[0] for r in d], ["SOS-26-01", "SOS-26-02", "SOS-26-03", "SOS-26-04", "Gözlem"])
            self.assertEqual(d[3][4], "Evet")

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_gizle_ve_cli(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "r.txt"
            y.write_text(METIN.replace("Kurgu Denetim Hizmetleri", "Denetçi Ayşe Deneme"), encoding="utf-8")
            agent.calistir(y, Path(tmp) / "d.xlsx", gizle=["Ayşe Deneme"], evet=True)
            self.assertNotIn("Ayşe", MESAJLAR[0])
            self.assertEqual(agent.main(["--termin", "x=1", "--evet"]), 2)
            self.assertEqual(agent.main(["--rapor", str(Path(tmp) / "yok.pdf"), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
