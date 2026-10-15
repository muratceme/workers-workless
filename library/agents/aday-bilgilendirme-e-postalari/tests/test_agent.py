"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import re
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

LISTE = KLASOR / "ornek_veri" / "adaylar.csv"
FIRMA = KLASOR / "ornek_veri" / "firma.txt"
BUGUN = date(2026, 10, 15)
MESAJLAR = []
DOGRU = {
    "S1": ("{POZISYON} pozisyonu için mülakat daveti",
           "Sayın {AD},\n\n{POZISYON} pozisyonu için yaptığınız başvuruyu değerlendirdik ve sizi {TARIH} tarihinde saat {SAAT} itibarıyla "
           "yüz yüze görüşmeye davet etmek istiyoruz.\nAdres: {YER}\nLütfen kimlik belgenizi yanınızda getiriniz ve katılımınızı bu e-postayı "
           "yanıtlayarak teyit ediniz.\n\nSaygılarımızla,\n{IMZA}"),
    "S2": ("{POZISYON} pozisyonu için online mülakat daveti",
           "Sayın {AD},\n\n{POZISYON} pozisyonu için sizi {TARIH} tarihinde saat {SAAT} itibarıyla online görüşmeye davet ediyoruz.\n"
           "Bağlantı: {YER}\nBağlantıyı görüşmeden önce denemenizi öneririz. Katılımınızı teyit etmenizi rica ederiz.\n\nSaygılarımızla,\n{IMZA}"),
    "S3": ("{POZISYON} başvurunuz hakkında",
           "Sayın {AD},\n\n{POZISYON} pozisyonuna gösterdiğiniz ilgi için teşekkür ederiz. Başvurunuzu dikkatle değerlendirdik; ancak bu süreçte "
           "başka adaylarla ilerlemeye karar verdik.\nKariyerinizde başarılar dileriz.\n\nSaygılarımızla,\n{IMZA}"),
    "S4": ("{POZISYON} mülakat süreci hakkında",
           "Sayın {AD},\n\n{POZISYON} pozisyonu için bize zaman ayırdığınız ve görüşmeye katıldığınız için teşekkür ederiz. Değerlendirmemiz "
           "sonucunda bu pozisyon için başka bir adayla ilerlemeye karar verdik.\nSorularınız için {ILETISIM} adresine yazabilirsiniz.\n\n"
           "Saygılarımızla,\n{IMZA}"),
    "S5": ("{SIRKET} — {POZISYON} iş teklifi",
           "Sayın {AD},\n\n{POZISYON} pozisyonu için size iş teklifimizi iletmekten memnuniyet duyuyoruz.\nÜcret: {UCRET}\n"
           "İşe başlama tarihi: {BASLANGIC}\nTeklifimize {SON_YANIT} tarihine kadar yanıt vermenizi rica ederiz. Resmî teklif mektubumuz "
           "ayrıca gönderilecektir.\n\nSaygılarımızla,\n{IMZA}"),
}


def sablonlar(dogru=True, sadece=None) -> dict:
    d = dict(DOGRU)
    if not dogru:
        d["S1"] = (d["S1"][0], d["S1"][1].replace("teyit ediniz.", "teyit ediniz. Görüşme 45 dakika sürecektir."))
        d["S3"] = (d["S3"][0], d["S3"][1].replace("başka adaylarla", "aradığımız yaş aralığı nedeniyle başka adaylarla"))
        d["S4"] = (d["S4"][0], d["S4"][1].replace("Sayın {AD},", "{AD}'a selamlar,").replace("ilerlemeye karar verdik.",
                                                                                         "ilerlemeye karar verdik. İleride mutlaka sizinle çalışacağız."))
        d["S5"] = (d["S5"][0], d["S5"][1].replace("{SON_YANIT} tarihine kadar", "en kısa sürede"))
    return {"sablonlar": [{"kimlik": k, "konu": v[0], "metin": v[1]} for k, v in d.items() if not sadece or k in sadece]}


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    ids = re.findall(r"^(S\d+) \|", kullanici, re.M)
    return sablonlar("<duzeltme>" in kullanici, ids)


class KodTesti(unittest.TestCase):
    def test_kombinasyon(self):
        a = agent.adaylari_oku(LISTE)
        k = agent.kombinasyonlar(a)
        self.assertEqual([(x["kimlik"], x["asama"], x["alt"]) for x in k],
                         [("S1", "Davet", "Yüz yüze"), ("S2", "Davet", "Online"), ("S3", "Olumsuz", "Başvuru değerlendirmesi"),
                          ("S4", "Olumsuz", "Mülakat sonrası"), ("S5", "Teklif", "")])
        self.assertEqual(a[4]["ucret"], "65.000 TL")
        self.assertEqual(agent.satir_sorunlari(a[3], BUGUN), ["E-posta adresi yok ya da geçersiz"])
        self.assertEqual(agent.satir_sorunlari(a[5], BUGUN), ["Mülakat tarihi geçmiş (12.10.2026)"])

    def test_denetim(self):
        k = agent.kombinasyonlar(agent.adaylari_oku(LISTE))
        self.assertEqual(agent.denetle(sablonlar(), k, 80), [])
        s = agent.denetle(sablonlar(False), k, 80)
        self.assertEqual(len(s), 5, s)
        self.assertTrue(any(x.startswith("S1 ") and "45" in x for x in s) and any(x.startswith("S3 ") and "ayrımcılık" in x for x in s)
                        and any(x.startswith("S4 ") and "ek getirilmiş" in x for x in s) and any(x.startswith("S4 ") and "söz" in x for x in s)
                        and any(x.startswith("S5 ") and "{SON_YANIT}" in x for x in s))
        self.assertTrue(any("konu satırı" in x for x in agent.denetle(sablonlar(), k, 30)))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "a.xlsx"
            s = agent.calistir(LISTE, FIRMA, cikti, evet=True, bugun=BUGUN)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            for gizli in ("Elif", "Koç", "example.com", "65000", "65.000", "22.10.2026", "10:30"):
                self.assertNotIn(gizli, MESAJLAR[0])
            self.assertNotIn("S2 |", MESAJLAR[1])  # düzeltmede yalnız sorunlu şablonlar
            self.assertEqual(len(s["epostalar"]), 5)
            self.assertEqual(sorted(x["ad"] for x, _ in s["atlanan"]), ["Kerem Ak", "Okan Er"])
            x, konu, metin = s["epostalar"][0]
            self.assertEqual(konu, "Muhasebe Uzmanı pozisyonu için mülakat daveti")
            self.assertIn("22.10.2026 tarihinde saat 10:30", metin)
            self.assertIn("Ücret: 65.000 TL", s["epostalar"][3][2])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Şablonlar", "E-postalar", "Gönderilemeyen"])
            self.assertEqual(s["csv"].read_text(encoding="utf-8-sig").count("@example.com"), 5)

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(LISTE, FIRMA, Path(tmp) / "a.xlsx", tur=0, evet=True, bugun=BUGUN)
            self.assertEqual((len(s["kalan"]), len(s["epostalar"]), len(s["atlanan"])), (5, 1, 6))
            self.assertEqual(agent.main(["--evet", "--bugun", "15.10.2026", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--evet", "--bugun", "dün"]), 1)


if __name__ == "__main__":
    unittest.main()
