"""
Aday Bilgilendirme E-postaları — Gerisi Bende AI Agent
İnsan Kaynakları › İnsan Kaynakları Uzman Yardımcısı

1. Kod aday listesini okur; her satırın aşamasını (mülakat daveti, olumsuz dönüş, teklif) ve gerekli bilgilerini
   (mülakat tarihi, saat, yer; teklif ücreti, başlangıç, yanıt son tarihi) denetler. Aşama + alt tür kombinasyonlarını çıkarır.
2. Model her kombinasyon için konu satırı ve yer tutuculu e-posta şablonu yazar. Aday adı, e-postası, ücreti gibi
   bilgiler modele gönderilmez; yalnız aşama, pozisyon adları ve şirketin üslup notu gönderilir.
3. Kod denetler:
   - Zorunlu yer tutucu eksik ya da bilinmeyen yer tutucu var mı; yer tutucuya ek getirilmiş mi ("{AD}'a")
   - Şablonda sabit sayı / tarih var mı (tarih, saat, ücret yer tutucuyla gelmeli)
   - Ayrımcılık sayılabilecek ifade (yaş, cinsiyet, medeni durum, hamilelik, din, engellilik, sağlık, uyruk …) var mı
   - Olumsuz dönüşte gerçekçi olmayan söz ("ileride mutlaka", "kesinlikle değerlendireceğiz") var mı
   - Konu satırı --konu-azami karakteri aşıyor mu
   Sorun varsa yalnız sorunlu şablonlar modele geri gönderilir (en fazla --tur kez).
4. Kod şablonları doldurur; e-postalar gönderilmez, Excel ve posta birleştirme için CSV yazılır.

Kullanım:
    python agent.py                                                     # örnek aday listesi
    python agent.py --girdi adaylar.xlsx --firma firma.txt --bugun 15.10.2026
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
YER_TUTUCULAR = ["{AD}", "{POZISYON}", "{SIRKET}", "{TARIH}", "{SAAT}", "{YER}", "{UCRET}", "{BASLANGIC}", "{SON_YANIT}", "{IMZA}", "{ILETISIM}"]
ZORUNLU = {"Davet": ["{AD}", "{POZISYON}", "{TARIH}", "{SAAT}", "{YER}", "{IMZA}"],
           "Olumsuz": ["{AD}", "{POZISYON}", "{IMZA}"],
           "Teklif": ["{AD}", "{POZISYON}", "{UCRET}", "{BASLANGIC}", "{SON_YANIT}", "{IMZA}"]}
GEREKLI_ALAN = {"Davet": (("tarih", "Mülakat Tarihi"), ("saat", "Mülakat Saati"), ("yer", "Mülakat Yeri / Bağlantı")),
                "Teklif": (("ucret", "Ücret"), ("baslangic", "Başlangıç Tarihi"), ("son_yanit", "Yanıt Son Tarihi"))}
AYRIMCI = re.compile(r"\b(yas(?:iniz|i|a|inizdan)?|cinsiyet\w*|medeni|evli\w*|bekar\w*|hamile\w*|gebe\w*|din(?:i|iniz)?|dini inanc\w*|mezhep\w*|etnik|irk\w*|"
                     r"engelli\w*|engel durum\w*|saglik durum\w*|uyruk\w*|askerlik\w*|cocuk sahibi)\b")
SOZ = re.compile(r"\b(ileride mutlaka|mutlaka (?:sizinle|size) (?:donecegiz|calisacagiz)|kesinlikle (?:degerlendirecegiz|donecegiz|arayacagiz)|"
                 r"ilk acilan pozisyon\w* (?:size|sizin))")
EPOSTA = re.compile(r"^[^@\s]+@[^@\s]+\.[a-z]{2,}$", re.I)
SEMA = {
    "type": "object",
    "properties": {"sablonlar": {"type": "array", "items": {"type": "object", "properties": {"kimlik": S, "konu": S, "metin": S},
                                                            "required": ["kimlik", "konu", "metin"], "additionalProperties": False}}},
    "required": ["sablonlar"],
    "additionalProperties": False,
}


def firma_oku(yol: Path) -> dict:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: firma bilgi dosyası bulunamadı.")
    metin = belge.metin_oku(yol)
    f = {"metin": metin}
    for alan, anahtar in (("sirket", "Şirket"), ("imza", "İmza"), ("iletisim", "İletişim")):
        m = re.search(rf"^\s*{anahtar}\s*:\s*(.+)$", metin, re.M | re.I)
        f[alan] = m.group(1).strip() if m else ""
    if not f["sirket"] or not f["imza"]:
        raise llm.LLMHatasi(f"{yol.name}: 'Şirket: …' ve 'İmza: …' satırları gerekli.")
    return f


def asama_norm(s: str) -> str:
    k = katla(s)
    return "Davet" if k.startswith(("davet", "mulakat")) else ("Olumsuz" if k.startswith(("olumsuz", "ret", "red")) else ("Teklif" if k.startswith("teklif") else ""))


def _tarih(v) -> str:
    t = tarih_coz(v)
    return f"{t:%d.%m.%Y}" if t else str(v or "").strip()


def _saat(v) -> str:
    if isinstance(v, (time, datetime)):
        return f"{v:%H:%M}"
    return str(v or "").strip()


def _ucret(v) -> str:
    s = str(v or "").strip()
    if isinstance(v, (int, float)) or re.fullmatch(r"\d+(?:[.,]\d+)?|\d{1,3}(?:\.\d{3})+", s):
        return f"{float(sayi_norm(s)):,.0f}".replace(",", ".") + " TL"
    return s


def adaylari_oku(yol: Path) -> list[dict]:
    r = kayitlar(yol, {"ad": ("Aday Adı", "Aday", "Ad Soyad"), "eposta": ("E-posta", "Eposta", "Mail"), "pozisyon": ("Pozisyon", "İlan"),
                       "asama": ("Aşama", "Durum", "E-posta Türü"), "tur": ("Mülakat Türü",), "tarih": ("Mülakat Tarihi",), "saat": ("Mülakat Saati", "Saat"),
                       "yer": ("Mülakat Yeri", "Yer / Bağlantı", "Bağlantı", "Yer"), "ucret": ("Ücret", "Brüt Ücret", "Teklif Ücreti"),
                       "baslangic": ("Başlangıç Tarihi", "Başlangıç"), "son_yanit": ("Yanıt Son Tarihi", "Son Yanıt")}, ("ad", "pozisyon", "asama"))
    out = []
    for x in r:
        if not str(x["ad"]).strip():
            continue
        asama = asama_norm(x["asama"])
        alt = ""
        if asama == "Davet":
            alt = "Online" if katla(x["tur"]).startswith(("online", "cevrim", "video")) else "Yüz yüze"
        elif asama == "Olumsuz":
            alt = "Mülakat sonrası" if str(x["tarih"]).strip() else "Başvuru değerlendirmesi"
        out.append({"ad": str(x["ad"]).strip(), "eposta": str(x["eposta"]).strip(), "pozisyon": str(x["pozisyon"]).strip(), "asama": asama, "alt": alt,
                    "tarih": _tarih(x["tarih"]), "tarih_d": tarih_coz(x["tarih"]), "saat": _saat(x["saat"]), "yer": str(x["yer"]).strip(),
                    "ucret": _ucret(x["ucret"]), "baslangic": _tarih(x["baslangic"]), "son_yanit": _tarih(x["son_yanit"])})
    return out


def kombinasyonlar(adaylar: list[dict]) -> list[dict]:
    out, gorulen = [], {}
    for x in adaylar:
        if not x["asama"]:
            continue
        a = (x["asama"], x["alt"])
        if a not in gorulen:
            gorulen[a] = {"kimlik": f"S{len(gorulen) + 1}", "asama": a[0], "alt": a[1], "pozisyonlar": []}
            out.append(gorulen[a])
        if x["pozisyon"] not in gorulen[a]["pozisyonlar"]:
            gorulen[a]["pozisyonlar"].append(x["pozisyon"])
    return out


def doldur(sablon: str, x: dict, f: dict) -> str:
    deger = {"{AD}": x["ad"], "{POZISYON}": x["pozisyon"], "{SIRKET}": f["sirket"], "{TARIH}": x["tarih"], "{SAAT}": x["saat"], "{YER}": x["yer"],
             "{UCRET}": x["ucret"], "{BASLANGIC}": x["baslangic"], "{SON_YANIT}": x["son_yanit"], "{IMZA}": f["imza"], "{ILETISIM}": f["iletisim"]}
    for k, v in deger.items():
        sablon = sablon.replace(k, v)
    return sablon.strip()


def denetle(t: dict, komb: list[dict], konu_azami: int) -> list[str]:
    sorun, sab = [], {s["kimlik"]: s for s in t["sablonlar"]}
    for k in komb:
        ad = f"{k['kimlik']} ({k['asama']} · {k['alt'] or '-'})"
        s = sab.get(k["kimlik"])
        if not s:
            sorun.append(f"{ad}: şablon yazılmadı")
            continue
        tum = s["konu"] + "\n" + s["metin"]
        eksik = [y for y in ZORUNLU[k["asama"]] if y not in tum]
        if eksik:
            sorun.append(f"{ad}: zorunlu yer tutucu eksik: {', '.join(eksik)}")
        bilinmeyen = sorted({y for y in re.findall(r"\{[^}]*\}", tum) if y not in YER_TUTUCULAR})
        if bilinmeyen:
            sorun.append(f"{ad}: bilinmeyen yer tutucu: {', '.join(bilinmeyen)}")
        ekli = re.findall(r"\{[A-Z_]+\}['’]?[a-zçğıöşü]+", tum)
        if ekli:
            sorun.append(f"{ad}: yer tutucuya ek getirilmiş ({', '.join(ekli)}); ada uymayabilir, cümleyi eksiz kurun")
        sade = re.sub(r"\{[^}]*\}", " ", tum)
        if sayilar(sade):
            sorun.append(f"{ad}: sabit sayı / tarih yazılmış ({', '.join(sorted(sayilar(sade)))}); tarih, saat ve ücret yer tutucuyla gelmeli")
        if AYRIMCI.search(katla(sade)):
            sorun.append(f"{ad}: ayrımcılık sayılabilecek ifade ('{AYRIMCI.search(katla(sade)).group(0)}')")
        if k["asama"] == "Olumsuz" and SOZ.search(katla(sade)):
            sorun.append(f"{ad}: gerçekçi olmayan söz ('{SOZ.search(katla(sade)).group(0)}')")
        if len(s["konu"]) > konu_azami:
            sorun.append(f"{ad}: konu satırı {len(s['konu'])} karakter, sınır {konu_azami}")
    return sorun


def uret(komb: list[dict], f: dict, konu_azami: int, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    sonuc, hedef = {}, komb
    if duzeltme is not None:  # düzeltmede yalnız sorunlu şablonlar yeniden yazılır
        sorunlu = {s.split(" ")[0] for s in duzeltme}
        sonuc.update({s["kimlik"]: s for s in (onceki or {}).get("sablonlar", []) if s["kimlik"] not in sorunlu})
        hedef = [k for k in komb if k["kimlik"] in sorunlu]
    satir = "\n".join(f"{k['kimlik']} | aşama: {k['asama']}{' · ' + k['alt'] if k['alt'] else ''} | zorunlu: {' '.join(ZORUNLU[k['asama']])} | "
                      f"pozisyonlar: {', '.join(k['pozisyonlar'])}" for k in hedef)
    ids = {k["kimlik"] for k in hedef}
    mesaj = (f"<firma_notu>\n{f['metin']}\n</firma_notu>\n<kombinasyonlar>\n{satir}\n</kombinasyonlar>\n<konu_azami>{konu_azami}</konu_azami>"
             + duzeltme_blogu(duzeltme, [s for s in (onceki or {}).get("sablonlar", []) if s["kimlik"] in ids]))
    for s in llm.json_iste(sistem, mesaj, SEMA).get("sablonlar", []):
        if s.get("kimlik") in ids:
            sonuc[s["kimlik"]] = s
    return {"sablonlar": [sonuc[k["kimlik"]] for k in komb if k["kimlik"] in sonuc]}


def satir_sorunlari(x: dict, bugun: date) -> list[str]:
    s = []
    if not x["asama"]:
        s.append("Aşama tanınmadı (Davet / Olumsuz / Teklif)")
    if not EPOSTA.match(x["eposta"]):
        s.append("E-posta adresi yok ya da geçersiz")
    for alan, ad in GEREKLI_ALAN.get(x["asama"], ()):
        if not x[alan]:
            s.append(f"{ad} boş")
    if x["asama"] == "Davet" and x["tarih_d"] and x["tarih_d"] < bugun:
        s.append(f"Mülakat tarihi geçmiş ({x['tarih']})")
    return s


def rapor_yaz(cikti: Path, komb, t, epostalar, atlanan, kalan, tur) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Şablonlar"
    sab = {s["kimlik"]: s for s in t["sablonlar"]}
    baslik(ws, ["Kimlik", "Aşama", "Alt Tür", "Pozisyonlar", "Konu", "Şablon", "Kontrol", "Onay"], (7, 10, 20, 30, 40, 80, 30, 8))
    for k in komb:
        sor = [s for s in kalan if s.startswith(k["kimlik"] + " ")]
        s = sab.get(k["kimlik"], {"konu": "", "metin": ""})
        ws.append([k["kimlik"], k["asama"], k["alt"], ", ".join(k["pozisyonlar"]), s["konu"], s["metin"], "; ".join(sor) or "Geçti", ""])
        boya(ws, [5, 6], AI_DOLGU)
        boya(ws, [7], KIRMIZI if sor else YESIL)
        boya(ws, [8], SARI)
    son_not(ws, "Önemli: E-postalar gönderilmedi. Şablonları İK sorumlusu onaylamalıdır; teklif e-postaları yazılı teklif / iş sözleşmesinin "
                "yerine geçmez.", kalan, tur)
    hizala(ws)
    m = wb.create_sheet("E-postalar")
    baslik(m, ["Aday", "E-posta", "Pozisyon", "Aşama", "Konu", "Metin"], (20, 28, 22, 10, 40, 90))
    for x, konu, metin in epostalar:
        m.append([x["ad"], x["eposta"], x["pozisyon"], x["asama"], konu, metin])
    hizala(m)
    a = wb.create_sheet("Gönderilemeyen")
    baslik(a, ["Aday", "E-posta", "Pozisyon", "Aşama", "Neden"], (20, 28, 22, 10, 60))
    for x, neden in atlanan:
        a.append([x["ad"], x["eposta"], x["pozisyon"], x["asama"], neden])
        boya(a, [5], KIRMIZI)
    hizala(a)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    yol = cikti.with_name(cikti.stem + "_birlestirme.csv")
    with yol.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["E-posta", "Konu", "Metin"])
        for x, konu, metin in epostalar:
            w.writerow([x["eposta"], konu, metin])
    return yol


def calistir(girdi: Path, firma_yolu: Path, cikti: Path, konu_azami: int = 80, tur: int = 2, evet: bool = False, bugun: date | None = None) -> dict:
    bugun = bugun or date.today()
    adaylar = adaylari_oku(girdi)
    f = firma_oku(firma_yolu)
    komb = kombinasyonlar(adaylar)
    if not komb:
        raise llm.LLMHatasi(f"{girdi.name}: tanınan aşama yok (Davet / Olumsuz / Teklif).")
    print(f"[OK] {len(adaylar)} aday · {len(komb)} şablon kombinasyonu")
    llm.onay_al(f"{len(komb)} aşama türü, pozisyon adları ve firma üslup notu gönderilecek (aday adı, e-posta, ücret ve tarih gönderilmez).", evet)
    t, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(komb, f, konu_azami, d, o), lambda x: denetle(x, komb, konu_azami), tur)
    sab = {s["kimlik"]: s for s in t["sablonlar"]}
    gecmeyen = {k["kimlik"] for k in komb if any(s.startswith(k["kimlik"] + " ") for s in kalan)}
    kimlik = {(k["asama"], k["alt"]): k["kimlik"] for k in komb}
    epostalar, atlanan = [], []
    for x in adaylar:
        sor = satir_sorunlari(x, bugun)
        kid = kimlik.get((x["asama"], x["alt"]))
        if not sor and (kid in gecmeyen or kid not in sab):
            sor = [f"Şablon {kid} kontrolden geçmedi"]
        if sor:
            atlanan.append((x, "; ".join(sor)))
        else:
            epostalar.append((x, doldur(sab[kid]["konu"], x, f), doldur(sab[kid]["metin"], x, f)))
    csv_yol = rapor_yaz(cikti, komb, t, epostalar, atlanan, kalan, yapilan)
    return {"adaylar": adaylar, "komb": komb, "taslak": t, "epostalar": epostalar, "atlanan": atlanan, "kalan": kalan, "tur": yapilan, "csv": csv_yol}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    p = argparse.ArgumentParser(description="Aday listesinden mülakat daveti, olumsuz dönüş ve teklif e-postaları hazırlar (göndermez).")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "adaylar.csv",
                   help="Aday listesi (.xlsx/.csv): Aday Adı, Pozisyon, Aşama; E-posta, Mülakat Türü / Tarihi / Saati / Yeri, Ücret, Başlangıç Tarihi, Yanıt Son Tarihi")
    p.add_argument("--firma", type=Path, default=BURASI / "ornek_veri" / "firma.txt", help="Firma bilgisi (.txt/.docx): 'Şirket:', 'İmza:', 'İletişim:' ve üslup notu")
    p.add_argument("--konu-azami", type=int, default=80, help="Konu satırı için en fazla karakter (varsayılan 80)")
    p.add_argument("--bugun", help="Geçmiş mülakat tarihi kontrolü için bugünün tarihi (GG.AA.YYYY; varsayılan: bugün)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "aday_epostalari.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        bugun = tarih_coz(a.bugun) if a.bugun else None
        if a.bugun and not bugun:
            raise llm.LLMHatasi(f"--bugun '{a.bugun}' tarih olarak okunamadı (GG.AA.YYYY).")
        s = calistir(a.girdi, a.firma, a.cikti, a.konu_azami, a.tur, a.evet, bugun)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(s['epostalar'])} e-posta hazır (gönderilmedi) · gönderilemeyen {len(s['atlanan'])}")
    for x, neden in s["atlanan"]:
        print(f"[!] {x['ad']} ({x['asama'] or '?'}): {neden}")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()} (+ {s['csv'].name})")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
