"""
Mağaza Ziyaret Raporu — Gerisi Bende AI Agent
Mağaza Operasyonları › Bölge Müdürü

1. Kod, ziyaret kontrol listesini okur (Bölüm, Madde, Sonuç: Uygun / Kısmen / Uygun değil / DD, Not), maddelere kimlik verir
   (K1, K2…) ve bölüm bazında uyum puanı hesaplar: Uygun 1, Kısmen 0,5, Uygun değil 0; DD hesaba girmez.
2. Model, kontrol listesi ve ziyaret notlarından rapor yazar: özet, güçlü yönler, aksiyon planı (ilgili madde kimlikleri,
   aksiyon, sorumlu rolü, termin, öncelik), mağaza müdürüyle görüşülenler ve takip ziyaretinde bakılacaklar.
3. Kod taslağı denetler: "Uygun değil" ve "Kısmen" her madde en az bir aksiyona bağlanmış mı; termin geçerli bir tarih mi ve
   ziyaret ile ziyaret + --azami gün arasında mı; güvenlik maddeleri için öncelik Yüksek mi; güçlü yön olarak yalnız "Uygun"
   maddeler mi gösterilmiş; kimlikler geçerli mi. Sorun varsa taslak modele geri gönderilir (en fazla --tur kez).
4. Çıktı: Excel (Rapor, Aksiyon Planı + 'Durum' sütunu, Bölüm Puanları, Kontrol Listesi) + Markdown rapor.

Kullanım:
    python agent.py                                              # örnek ziyaretle dener
    python agent.py --kontrol kontrol.xlsx --notlar notlar.docx --tarih 14.10.2026
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

import belge
import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
PUAN = {"Uygun": 1.0, "Kısmen": 0.5, "Uygun değil": 0.0}
GUVENLIK = ("guvenlik", "isg", "acil", "yangin", "tahliye", "elektrik", "is sagligi")

S = {"type": "string"}
SEMA = {
    "type": "object",
    "properties": {
        "ozet": S,
        "guclu_yonler": {"type": "array", "items": {"type": "object", "properties": {"metin": S, "maddeler": {"type": "array", "items": S}},
                                                     "required": ["metin", "maddeler"], "additionalProperties": False}},
        "aksiyonlar": {"type": "array", "items": {"type": "object", "properties": {
            "maddeler": {"type": "array", "items": S}, "aksiyon": S, "sorumlu": S, "termin": S, "oncelik": {"type": "string", "enum": ["Yüksek", "Orta", "Düşük"]}},
            "required": ["maddeler", "aksiyon", "sorumlu", "termin", "oncelik"], "additionalProperties": False}},
        "gorusulenler": {"type": "array", "items": S},
        "takip_konulari": {"type": "array", "items": S},
    },
    "required": ["ozet", "guclu_yonler", "aksiyonlar", "gorusulenler", "takip_konulari"],
    "additionalProperties": False,
}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def sonuc(x) -> str | None:
    k = katla(x)
    if k in ("uygun", "evet", "var", "tamam", "ok", "2"):
        return "Uygun"
    if k in ("kismen", "kismi", "1"):
        return "Kısmen"
    if k in ("uygun degil", "hayir", "yok", "eksik", "0"):
        return "Uygun değil"
    if k in ("dd", "uygulanamaz", "kapsam disi", "na", "n a"):
        return "DD"
    return None


def _tablo(yol: Path) -> list[list]:
    if yol.suffix.lower() in (".xlsx", ".xlsm"):
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            return [list(r) for r in wb.active.iter_rows(values_only=True) if any(c not in (None, "") for c in r)]
        finally:
            wb.close()
    metin = belge.txt_oku(yol)
    ayirac = max(";\t,", key=metin.splitlines()[0].count) if metin.strip() else ";"
    return [r for r in csv.reader(metin.splitlines(), delimiter=ayirac) if any(c.strip() for c in r)]


def kontrol_oku(yol: Path) -> tuple[list[dict], list[str]]:
    s = _tablo(yol)
    if not s:
        raise belge.BelgeHatasi(f"{yol.name}: kontrol listesi boş.")
    b = [katla(c) for c in s[0]]
    try:
        jb, jm, js = b.index("bolum"), b.index("madde"), b.index("sonuc")
    except ValueError as h:
        raise belge.BelgeHatasi(f"{yol.name}: 'Bölüm', 'Madde' ve 'Sonuç' sütunları gerekli.") from h
    jn = b.index("not") if "not" in b else None
    maddeler, hatalar = [], []
    for n, r in enumerate(s[1:], 2):
        al = lambda j: str(r[j]).strip() if j is not None and j < len(r) and r[j] is not None else ""
        if not al(jm):
            continue
        sn = sonuc(al(js))
        if sn is None:
            hatalar.append(f"Satır {n}: sonuç '{al(js)}' tanınmadı (Uygun / Kısmen / Uygun değil / DD); madde alınmadı.")
            continue
        maddeler.append({"id": f"K{len(maddeler) + 1}", "bolum": al(jb) or "Genel", "madde": al(jm), "sonuc": sn, "not": al(jn)})
    if not maddeler:
        raise belge.BelgeHatasi(f"{yol.name}: okunabilen madde yok.")
    return maddeler, hatalar


def puanla(maddeler: list[dict]) -> list[dict]:
    bolumler = {}
    for m in maddeler:
        bolumler.setdefault(m["bolum"], []).append(m)
    out = []
    for ad, l in list(bolumler.items()) + [("TOPLAM", maddeler)]:
        u = [m for m in l if m["sonuc"] != "DD"]
        out.append({"bolum": ad, "madde": len(l), "uygun": sum(m["sonuc"] == "Uygun" for m in l), "kismen": sum(m["sonuc"] == "Kısmen" for m in l),
                    "degil": sum(m["sonuc"] == "Uygun değil" for m in l), "puan": round(sum(PUAN[m["sonuc"]] for m in u) * 100 / len(u), 1) if u else None})
    return out


def _tarih(x) -> date | None:
    try:
        return datetime.strptime(str(x).strip()[:10], "%d.%m.%Y").date()
    except ValueError:
        return None


def guvenlik_mi(m: dict) -> bool:
    return any(k in katla(m["bolum"] + " " + m["madde"]) for k in GUVENLIK)


def denetle(t: dict, maddeler: list[dict], ziyaret: date, azami: int = 90) -> list[str]:
    idler = {m["id"]: m for m in maddeler}
    sorunlar = []
    kapsanan = {}
    for a in t["aksiyonlar"]:
        ad = a["aksiyon"][:60]
        gecersiz = [i for i in a["maddeler"] if i not in idler]
        if gecersiz:
            sorunlar.append(f"Aksiyon '{ad}': geçersiz madde kimliği {gecersiz}")
        g = _tarih(a["termin"])
        if g is None:
            sorunlar.append(f"Aksiyon '{ad}': termin '{a['termin']}' GG.AA.YYYY biçiminde değil")
        elif not ziyaret <= g <= ziyaret + timedelta(days=azami):
            sorunlar.append(f"Aksiyon '{ad}': termin {a['termin']} ziyaret ({ziyaret:%d.%m.%Y}) ile {azami} gün sonrası arasında değil")
        if not a["sorumlu"].strip():
            sorunlar.append(f"Aksiyon '{ad}': sorumlu yazılmamış")
        for i in a["maddeler"]:
            if i in idler:
                kapsanan.setdefault(i, []).append(a)
    for m in maddeler:
        if m["sonuc"] in ("Uygun değil", "Kısmen") and m["id"] not in kapsanan:
            sorunlar.append(f"{m['id']} ({m['sonuc']}: {m['madde']}) için aksiyon yok")
        if m["sonuc"] == "Uygun değil" and guvenlik_mi(m) and m["id"] in kapsanan and not any(a["oncelik"] == "Yüksek" for a in kapsanan[m["id"]]):
            sorunlar.append(f"{m['id']} güvenlik maddesi ({m['madde']}) uygun değil; aksiyon önceliği Yüksek olmalı")
    for g in t["guclu_yonler"]:
        for i in g["maddeler"]:
            if i not in idler:
                sorunlar.append(f"Güçlü yön '{g['metin'][:50]}': geçersiz madde kimliği {i}")
            elif idler[i]["sonuc"] != "Uygun":
                sorunlar.append(f"Güçlü yön '{g['metin'][:50]}': {i} '{idler[i]['sonuc']}' olduğu hâlde güçlü yön gösterilmiş")
    return sorunlar


def uret(maddeler: list[dict], puanlar: list[dict], notlar: str, ziyaret: date, azami: int, duzeltme: list[str] | None = None, onceki: dict | None = None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    liste = "\n".join(f"{m['id']} | {m['bolum']} | {m['madde']} | {m['sonuc']}" + (f" | not: {llm.maskele(m['not'])}" if m["not"] else "") for m in maddeler)
    pn = "\n".join(f"{p['bolum']}: %{p['puan']}" for p in puanlar)
    parcalar = [f"<ziyaret_tarihi>{ziyaret:%d.%m.%Y} (termin en geç {ziyaret + timedelta(days=azami):%d.%m.%Y})</ziyaret_tarihi>", f"<kontrol_listesi>\n{liste}\n</kontrol_listesi>",
                f"<bolum_puanlari>\n{pn}\n</bolum_puanlari>", f"<ziyaret_notlari>\n{llm.maskele(notlar)}\n</ziyaret_notlari>"]
    if duzeltme:
        parcalar.append("<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False) + "\nKodun bulduğu sorunlar:\n" + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")
    return llm.json_iste(sistem, "\n".join(parcalar), SEMA)


BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for j, w in enumerate(genislik):
        ws.column_dimensions[chr(65 + j)].width = w


def rapor_metni(t: dict, puanlar: list[dict], ziyaret: date, idler: dict) -> str:
    top = puanlar[-1]["puan"]
    out = [f"# Mağaza Ziyaret Raporu · {ziyaret:%d.%m.%Y}", f"Genel uyum puanı: %{top}", "", t["ozet"], "", "## Bölüm puanları"]
    out += [f"- {p['bolum']}: %{p['puan']}" for p in puanlar[:-1]]
    out += ["", "## Güçlü yönler"] + [f"- {g['metin']}" for g in t["guclu_yonler"]]
    out += ["", "## Aksiyon planı"] + [f"- [{a['oncelik']}] {a['aksiyon']} — {a['sorumlu']}, termin {a['termin']} ({', '.join(a['maddeler'])})" for a in t["aksiyonlar"]]
    out += ["", "## Mağaza müdürüyle görüşülenler"] + [f"- {g}" for g in t["gorusulenler"]]
    out += ["", "## Takip ziyaretinde bakılacaklar"] + [f"- {g}" for g in t["takip_konulari"]]
    return "\n".join(out)


def rapor_yaz(cikti: Path, maddeler: list[dict], puanlar: list[dict], t: dict, kalan: list[str], tur: int, ziyaret: date) -> Path:
    idler = {m["id"]: m for m in maddeler}
    wb = Workbook()
    ws = wb.active
    ws.title = "Rapor"
    ws.column_dimensions["A"].width = 120
    ws.append([f"Mağaza ziyaret raporu · {ziyaret:%d.%m.%Y} · genel uyum %{puanlar[-1]['puan']}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([rapor_metni(t, puanlar, ziyaret, idler)])
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Uygun olmayan her maddeye terminli aksiyon bağlandı")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append(["Önemli: Yapay zekâ taslağıdır; bölge müdürü kontrol edip mağaza müdürüyle paylaşmalıdır."])
    for satir in ws.iter_rows():
        for h in satir:
            h.alignment = UST
    a = wb.create_sheet("Aksiyon Planı")
    _baslik(a, ["Öncelik", "Aksiyon", "İlgili Madde", "Sorumlu", "Termin", "Durum", "Kapanış Notu"], (9, 60, 40, 18, 11, 12, 30))
    for x in sorted(t["aksiyonlar"], key=lambda x: (["Yüksek", "Orta", "Düşük"].index(x["oncelik"]), _tarih(x["termin"]) or date.max)):
        a.append([x["oncelik"], x["aksiyon"], "; ".join(f"{i} {idler[i]['madde']}" for i in x["maddeler"] if i in idler), x["sorumlu"], x["termin"], "Açık", ""])
        a.cell(a.max_row, 1).fill = KIRMIZI if x["oncelik"] == "Yüksek" else (SARI if x["oncelik"] == "Orta" else YESIL)
        a.cell(a.max_row, 6).fill = SARI
        for h in a[a.max_row]:
            h.alignment = UST
    p = wb.create_sheet("Bölüm Puanları")
    _baslik(p, ["Bölüm", "Madde", "Uygun", "Kısmen", "Uygun Değil", "Uyum %"], (24, 8, 8, 8, 11, 9))
    for x in puanlar:
        p.append([x["bolum"], x["madde"], x["uygun"], x["kismen"], x["degil"], x["puan"]])
        if x["puan"] is not None:
            p.cell(p.max_row, 6).fill = YESIL if x["puan"] >= 90 else (SARI if x["puan"] >= 70 else KIRMIZI)
    k = wb.create_sheet("Kontrol Listesi")
    _baslik(k, ["Kimlik", "Bölüm", "Madde", "Sonuç", "Not"], (8, 22, 50, 12, 50))
    renk = {"Uygun": YESIL, "Kısmen": SARI, "Uygun değil": KIRMIZI}
    for m in maddeler:
        k.append([m["id"], m["bolum"], m["madde"], m["sonuc"], m["not"]])
        if m["sonuc"] in renk:
            k.cell(k.max_row, 4).fill = renk[m["sonuc"]]
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    md = cikti.with_suffix(".md")
    md.write_text(rapor_metni(t, puanlar, ziyaret, idler) + "\n", encoding="utf-8")
    return md


def ziyaret_tarihi(notlar: str) -> date | None:
    m = re.search(r"(\d{2}\.\d{2}\.\d{4})", notlar)
    return _tarih(m.group(1)) if m else None


def calistir(kontrol: Path, notlar_yol: Path | None, cikti: Path, ziyaret: date | None = None, tur: int = 2, azami: int = 90, evet: bool = False) -> dict:
    maddeler, hatalar = kontrol_oku(kontrol)
    notlar = belge.metin_oku(notlar_yol) if notlar_yol else ""
    ziyaret = ziyaret or ziyaret_tarihi(notlar)
    if ziyaret is None:
        raise belge.BelgeHatasi("Ziyaret tarihi bulunamadı; --tarih GG.AA.YYYY verin.")
    puanlar = puanla(maddeler)
    for h in hatalar:
        print(f"[!] {h}")
    print(f"[OK] {len(maddeler)} madde · genel uyum %{puanlar[-1]['puan']} · uygun değil {puanlar[-1]['degil']} · kısmen {puanlar[-1]['kismen']}")
    llm.onay_al("Kontrol listesi, bölüm puanları ve ziyaret notları (telefon / e-posta maskeli) rapor yazımı için gönderilecek.", evet)
    taslak = uret(maddeler, puanlar, notlar, ziyaret, azami)
    kalan = denetle(taslak, maddeler, ziyaret, azami)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(maddeler, puanlar, notlar, ziyaret, azami, kalan, taslak)
        kalan = denetle(taslak, maddeler, ziyaret, azami)
    md = rapor_yaz(cikti, maddeler, puanlar, taslak, kalan, yapilan, ziyaret)
    return {"maddeler": maddeler, "puanlar": puanlar, "taslak": taslak, "kalan": kalan, "tur": yapilan, "md": md, "hatalar": hatalar}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Mağaza ziyaret kontrol listesi ve notlarından aksiyon ve takip raporu hazırlar.")
    p.add_argument("--kontrol", type=Path, default=BURASI / "ornek_veri" / "kontrol_listesi.csv", help="Kontrol listesi (.csv / .xlsx): Bölüm, Madde, Sonuç, Not")
    p.add_argument("--notlar", type=Path, help="Ziyaret notları (.txt, .md, .docx, .pdf)")
    p.add_argument("--tarih", help="Ziyaret tarihi GG.AA.YYYY (verilmezse notlardaki ilk tarih)")
    p.add_argument("--azami", type=int, default=90, help="Aksiyon termini en fazla ziyaretten kaç gün sonra olabilir (varsayılan 90)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "magaza_ziyaret_raporu.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    notlar = a.notlar or (BURASI / "ornek_veri" / "ziyaret_notlari.txt" if a.kontrol == BURASI / "ornek_veri" / "kontrol_listesi.csv" else None)
    zt = _tarih(a.tarih) if a.tarih else None
    if (a.tarih and zt is None) or a.azami < 1:
        print("[X] --tarih GG.AA.YYYY, --azami en az 1 olmalı")
        return 2
    try:
        s = calistir(a.kontrol, notlar, a.cikti, zt, a.tur, a.azami, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] Rapor hazır · {len(s['taslak']['aksiyonlar'])} aksiyon · güçlü yön {len(s['taslak']['guclu_yonler'])}")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()} (+ {s['md'].name})")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
