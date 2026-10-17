"""
Rapor Yorum ve Aksiyon Notu — Gerisi Bende AI Agent
Yönetim raporları (çok sektörlü)

Bir Excel raporundan (sitedeki kod bloklarının çıktıları ya da kendi raporunuz) yönetici için kısa bir yorum notu yazar:
özet, öne çıkan bulgular ve aksiyon önerileri.
1. Kod raporun sayfalarını okur; hücreleri Excel'de göründükleri gibi biçimler (%55,6 · 967.029,08). Uzun tablolarda ilk
   satırlar ile her sayısal sütunun toplamı, ortalaması, en büyüğü ve en küçüğü (hangi satırda olduğuyla) gönderilir.
   İstenen sütunlardaki adlar (ör. temsilci, müşteri) takma adla gönderilir ve notta geri açılır; telefon, e-posta, TCKN,
   IBAN maskelenir.
2. Model raporu yorumlar: en fazla 5 maddelik özet, önem dereceli bulgular (her biri bir sayfaya dayanır) ve bulgulara
   bağlı aksiyon önerileri.
3. Kod notu denetler: nottaki her sayı raporda (ya da kodun hesapladığı özetlerde) var mı (yuvarlama payıyla), bulgunun
   sayfası raporda var mı, aksiyon geçerli bir bulguya mı bağlı. Sorun varsa not modele geri gönderilir (en fazla --tur kez).

Kullanım:
    python agent.py                                                    # örnek: satış hedef-gerçekleşme raporu
    python agent.py --rapor rapor.xlsx --okuyucu "Bölge Satış Müdürü" --odak "hedefin gerisinde kalanlar"
    python agent.py --rapor gecikme.xlsx --takma-ad Müşteri
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import llm

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
RENK = {"Kritik": "FDE2E1", "Önemli": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}
TAM_SATIR = 40          # bu kadar satıra kadar tablo tamamen gönderilir
ILK_SATIR = 25          # daha uzun tablolarda gönderilen ilk satır sayısı
AZAMI_KARAKTER = 60000


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def tr_sayi(v: float | Decimal, ondalik: int) -> str:
    s = f"{float(v):,.{ondalik}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def bicimle(v, fmt: str) -> str:
    """Hücreyi Excel'de göründüğüne yakın yazar: yüzde biçimi → %55,6; ondalıklı → 2 hane; tam sayı → binlik ayırıcı."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Evet" if v else "Hayır"
    if isinstance(v, datetime):
        return f"{v:%d.%m.%Y}" if (v.hour, v.minute) == (0, 0) else f"{v:%d.%m.%Y %H:%M}"
    if isinstance(v, date):
        return f"{v:%d.%m.%Y}"
    if isinstance(v, (int, float)):
        if "%" in (fmt or ""):
            return ("-%" if v < 0 else "%") + tr_sayi(abs(v) * 100, 1)
        if float(v).is_integer():
            return tr_sayi(v, 0)
        return tr_sayi(v, 2)
    return " ".join(str(v).split())


def para(x) -> Decimal | None:
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1 or (s.count(".") == 1 and len(s.split(".")[1]) == 3):
        s = s.replace(".", "")
    try:
        return Decimal(s) if s not in ("", "-", ".") else None
    except InvalidOperation:
        return None


SAYI_DESEN = re.compile(r"(?<![\w.,])-?\d+(?:[.,]\d+)*")


def sayi_tokenlari(metin: str) -> list[tuple[Decimal, int]]:
    """Metindeki sayılar ve yazıldıkları ondalık hane sayısı. Tarih ve saat parçaları ayrı sayılar olarak alınır."""
    out = []
    for m in re.finditer(r"\d{1,2}[./]\d{1,2}[./]\d{2,4}|\d{1,2}:\d{2}", metin):
        out += [(Decimal(x), 0) for x in re.findall(r"\d+", m.group(0))]
    temiz = re.sub(r"\d{1,2}[./]\d{1,2}[./]\d{2,4}|\d{1,2}:\d{2}", " ", metin)
    for x in SAYI_DESEN.findall(temiz):
        v = para(x)
        if v is None:
            continue
        ondalik = 0
        t = x.lstrip("-")
        if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", t) or re.fullmatch(r"\d+,\d+", t):
            ondalik = len(t.split(",")[1]) if "," in t else 0
        elif re.fullmatch(r"\d+\.\d{1,2}", t):
            ondalik = len(t.split(".")[1])
        out.append((abs(v), ondalik))
    return out


def eslesir(v: Decimal, ondalik: int, izinli: set[Decimal]) -> bool:
    if v in izinli:
        return True
    q = Decimal(1).scaleb(-ondalik)
    return any(abs(x).quantize(q, rounding=ROUND_HALF_UP) == v for x in izinli)


# ----------------------------------------------------------------------------
# Rapor okuma
# ----------------------------------------------------------------------------

def sayfalari_oku(yol: Path, takma_sutunlar: list[str]) -> tuple[list[dict], dict[str, str]]:
    """Her sayfa: ad, satırlar (biçimlenmiş metin), sayısal sütun özetleri. Takma ad haritası: gerçek → takma."""
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: rapor bulunamadı.")
    if yol.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise llm.LLMHatasi(f"{yol.name}: yalnız .xlsx / .xlsm raporlar okunur.")
    wb = load_workbook(yol, data_only=True)
    hedef = {katla(x) for x in takma_sutunlar}
    takma: dict[str, str] = {}
    sayfalar = []
    for ws in wb.worksheets:
        satirlar = [[(c.value, c.number_format) for c in r] for r in ws.iter_rows()]
        while satirlar and all(v in (None, "") for v, _ in satirlar[-1]):
            satirlar.pop()
        if not satirlar:
            continue
        genislik = max(max((i + 1 for i, (v, _) in enumerate(r) if v not in (None, "")), default=0) for r in satirlar)
        satirlar = [r[:genislik] for r in satirlar]
        # başlık satırı: en çok metin hücresi olan ilk 10 satırdan biri
        bas = max(range(min(10, len(satirlar))), key=lambda i: (sum(isinstance(v, str) for v, _ in satirlar[i]), -i))
        basliklar = [katla(v) for v, _ in satirlar[bas]]
        gizli_sutun = {j for j, b in enumerate(basliklar) if b in hedef}
        metin = []
        for i, r in enumerate(satirlar):
            hucre = []
            for j, (v, f) in enumerate(r):
                s = bicimle(v, f)
                if i > bas and j in gizli_sutun and s:
                    takma.setdefault(s, f"Kişi-{len(takma) + 1}")
                    s = takma[s]
                hucre.append(s)
            metin.append(hucre)
        ozet = []
        govde = satirlar[bas + 1:]
        if len(govde) > TAM_SATIR:
            etiket_j = next((j for j in range(genislik) if all(isinstance(r[j][0], (str, type(None))) for r in govde[:20])), 0)
            for j in range(genislik):
                deger = [(r[j][0], i) for i, r in enumerate(govde) if isinstance(r[j][0], (int, float)) and not isinstance(r[j][0], bool)]
                if len(deger) < len(govde) / 2:
                    continue
                f = govde[0][j][1]
                vs = [v for v, _ in deger]
                buyuk, kucuk = max(deger), min(deger)
                etk = lambda k: metin[bas + 1 + k][etiket_j] or f"satır {bas + 2 + k}"  # noqa: E731
                ozet.append(f"{metin[bas][j] or get_column_letter(j + 1)}: toplam {bicimle(sum(vs), f)} · ortalama {bicimle(sum(vs) / len(vs), f)} · "
                            f"en büyük {bicimle(buyuk[0], f)} ({etk(buyuk[1])}) · en küçük {bicimle(kucuk[0], f)} ({etk(kucuk[1])}) · {len(vs)} değer")
            metin_gonder = metin[:bas + 1 + ILK_SATIR]
            kesik = len(govde) - ILK_SATIR
        else:
            metin_gonder, kesik = metin, 0
        sayfalar.append({"ad": ws.title, "satirlar": metin_gonder, "tum": metin, "ozet": ozet, "kesik": kesik, "baslik_satiri": bas})
    wb.close()
    if not sayfalar:
        raise llm.LLMHatasi(f"{yol.name}: dolu sayfa yok.")
    return sayfalar, takma


def rapor_metni(sayfalar: list[dict]) -> str:
    parca = []
    for s in sayfalar:
        l = [f'<sayfa ad="{s["ad"]}">']
        for r in s["satirlar"]:
            if any(r):
                l.append(" | ".join(r).rstrip(" |"))
        if s["kesik"]:
            l.append(f"(… {s['kesik']} satır daha gönderilmedi; aşağıda bütün satırların kod tarafından hesaplanan özeti var)")
        if s["ozet"]:
            l.append("Kodun hesapladığı sütun özetleri (bütün satırlar):")
            l += [f"- {x}" for x in s["ozet"]]
        l.append("</sayfa>")
        parca.append("\n".join(l))
    return llm.maskele("\n".join(parca))


def izinli_sayilar(metin: str) -> set[Decimal]:
    return {v for v, _ in sayi_tokenlari(metin)}


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

SEMA = {"type": "object", "properties": {
    "baslik": S,
    "ozet": {"type": "array", "items": S},
    "bulgular": {"type": "array", "items": {"type": "object", "properties": {
        "metin": S, "sayfa": S, "onem": {"type": "string", "enum": ["Kritik", "Önemli", "Bilgi"]}},
        "required": ["metin", "sayfa", "onem"], "additionalProperties": False}},
    "aksiyonlar": {"type": "array", "items": {"type": "object", "properties": {
        "is": S, "sorumlu": S, "bulgu_no": {"type": "integer"}}, "required": ["is", "sorumlu", "bulgu_no"], "additionalProperties": False}},
    "veri_notlari": {"type": "array", "items": S}},
    "required": ["baslik", "ozet", "bulgular", "aksiyonlar", "veri_notlari"], "additionalProperties": False}


def metinler(n: dict) -> list[tuple[str, str]]:
    out = [("Başlık", n["baslik"])] + [(f"Özet {i}", x) for i, x in enumerate(n["ozet"], 1)]
    out += [(f"Bulgu {i}", b["metin"]) for i, b in enumerate(n["bulgular"], 1)]
    out += [(f"Aksiyon {i}", a["is"]) for i, a in enumerate(n["aksiyonlar"], 1)] + [(f"Veri notu {i}", x) for i, x in enumerate(n["veri_notlari"], 1)]
    return out


def denetle(n: dict, izinli: set[Decimal], sayfa_adlari: set[str]) -> list[str]:
    sorun = []
    for ad, m in metinler(n):
        fazla = [v for v, d in sayi_tokenlari(m) if not eslesir(v, d, izinli)]
        if fazla:
            sorun.append(f"{ad}: raporda olmayan sayı: {', '.join(tr_sayi(v, max(0, -v.as_tuple().exponent)) for v in fazla)} "
                         "(rapordaki rakamı aynen ya da daha az ondalıkla yazın; oran, fark veya pay hesaplamayın)")
    if not 1 <= len(n["ozet"]) <= 5:
        sorun.append(f"Özet: {len(n['ozet'])} madde; 1 ile 5 arasında olmalı")
    if not n["bulgular"]:
        sorun.append("Bulgular: hiç bulgu yok")
    if len(n["bulgular"]) > 8:
        sorun.append(f"Bulgular: {len(n['bulgular'])} bulgu; en fazla 8")
    for i, b in enumerate(n["bulgular"], 1):
        if b["sayfa"] not in sayfa_adlari:
            sorun.append(f"Bulgu {i}: sayfa '{b['sayfa']}' raporda yok ({', '.join(sorted(sayfa_adlari))})")
    for i, a in enumerate(n["aksiyonlar"], 1):
        if not 1 <= a["bulgu_no"] <= len(n["bulgular"]):
            sorun.append(f"Aksiyon {i}: bulgu_no {a['bulgu_no']} geçerli bir bulgu değil (1-{len(n['bulgular'])})")
    for ad, m in metinler(n):
        if ek := re.search(r"\bKişi-\d+['’]\w+", m):
            sorun.append(f"{ad}: takma ada ek getirilmiş ('{ek.group(0)}'); gerçek ad geri konunca ek yanlış olur. Eksiz yazın "
                         "(ör. 'Kişi-4 adlı temsilcinin', 'Kişi-4 için', 'Kişi-4: …')")
        if re.search(r"\[(TCKN|IBAN|TELEFON|E-POSTA|LINKEDIN)\]", m):
            sorun.append(f"{ad}: maskeli yer tutucu kullanılmış; o bilgiyi yazmayın")
    return sorun


def uret(rapor: str, okuyucu: str, odak: str, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    mesaj = f"<okuyucu>{okuyucu}</okuyucu>\n<odak>{odak or 'genel değerlendirme'}</odak>\n<rapor>\n{rapor}\n</rapor>"
    if duzeltme:
        mesaj += ("\n<duzeltme>\nÖnceki not: " + json.dumps(onceki, ensure_ascii=False) + "\nKodun bulduğu sorunlar:\n"
                  + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")
    return llm.json_iste(sistem, mesaj, SEMA)


def duzeltme_dongusu(uret_f, denetle_f, tur: int) -> tuple[dict, list[str], int]:
    taslak = uret_f(None, None)
    kalan = denetle_f(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret_f(kalan, taslak)
        kalan = denetle_f(taslak)
    return taslak, kalan, yapilan


def geri_ac(n: dict, takma: dict[str, str]) -> dict:
    ters = sorted(((t, g) for g, t in takma.items()), key=lambda x: -len(x[0]))

    def ac(s):
        if isinstance(s, str):
            for t, g in ters:
                s = re.sub(rf"\b{re.escape(t)}\b", g, s)
            return s
        if isinstance(s, list):
            return [ac(x) for x in s]
        if isinstance(s, dict):
            return {k: ac(v) for k, v in s.items()}
        return s
    return ac(n)


# ----------------------------------------------------------------------------
# Çıktı
# ----------------------------------------------------------------------------

def not_metni(n: dict, rapor_adi: str, okuyucu: str) -> str:
    l = [n["baslik"], f"Kaynak: {rapor_adi} · Okuyucu: {okuyucu} · {datetime.now():%d.%m.%Y}", "", "ÖZET"]
    l += [f"• {x}" for x in n["ozet"]] + ["", "BULGULAR"]
    l += [f"{i}. [{b['onem']}] {b['metin']} (sayfa: {b['sayfa']})" for i, b in enumerate(n["bulgular"], 1)]
    l += ["", "AKSİYON ÖNERİLERİ"] + [f"{i}. {a['is']} — Sorumlu: {a['sorumlu']} (bulgu {a['bulgu_no']})" for i, a in enumerate(n["aksiyonlar"], 1)]
    if n["veri_notlari"]:
        l += ["", "VERİ NOTLARI"] + [f"• {x}" for x in n["veri_notlari"]]
    l += ["", "Bu not yapay zekâ ile hazırlanmış bir taslaktır; rakamlar kaynak raporla kod tarafından karşılaştırılmıştır, "
              "yorum ve öneriler karar vermeden önce kontrol edilmelidir."]
    return "\n".join(l)


def _baslik(ws, basliklar, genislik) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font, h.alignment = BASLIK_DOLGU, BASLIK_YAZI, UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def rapor_yaz(cikti: Path, n: dict, kalan: list[str], tur: int, rapor_adi: str, okuyucu: str) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Not"
    ws.column_dimensions["A"].width = 120
    for s in not_metni(n, rapor_adi, okuyucu).splitlines():
        ws.append([s])
        ws.cell(ws.max_row, 1).alignment = UST
        if s in ("ÖZET", "BULGULAR", "AKSİYON ÖNERİLERİ", "VERİ NOTLARI") or ws.max_row == 1:
            ws.cell(ws.max_row, 1).font = Font(bold=True)
    b = wb.create_sheet("Bulgular")
    _baslik(b, ["No", "Önem", "Bulgu", "Sayfa"], (5, 9, 100, 24))
    for i, x in enumerate(n["bulgular"], 1):
        b.append([i, x["onem"], x["metin"], x["sayfa"]])
        b.cell(b.max_row, 2).fill = PatternFill("solid", fgColor=RENK[x["onem"]])
        b.cell(b.max_row, 3).fill = AI_DOLGU
        b.cell(b.max_row, 3).alignment = UST
    a = wb.create_sheet("Aksiyonlar")
    _baslik(a, ["No", "İş", "Sorumlu", "Bulgu", "Durum", "Termin"], (5, 80, 24, 7, 12, 12))
    for i, x in enumerate(n["aksiyonlar"], 1):
        a.append([i, x["is"], x["sorumlu"], x["bulgu_no"], "Açık", None])
        a.cell(a.max_row, 2).fill = AI_DOLGU
        a.cell(a.max_row, 2).alignment = UST
    k = wb.create_sheet("Kontrol")
    k.append([f"Kaynak rapor: {rapor_adi} · Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    k.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Nottaki bütün sayılar raporda var; kod denetiminden sorunsuz geçti")])
    k.cell(2, 1).fill = KIRMIZI if kalan else YESIL
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    txt = cikti.with_suffix(".txt")
    txt.write_text(not_metni(n, rapor_adi, okuyucu), encoding="utf-8")
    return txt


def calistir(rapor: Path, cikti: Path, okuyucu: str = "Yönetici", odak: str = "", takma_ad: list[str] | None = None, tur: int = 2,
             evet: bool = False) -> dict:
    sayfalar, takma = sayfalari_oku(rapor, takma_ad or [])
    metin = rapor_metni(sayfalar)
    if len(metin) > AZAMI_KARAKTER:
        raise llm.LLMHatasi(f"Rapor çok büyük ({len(metin):,} karakter). Yalnız yorumlanacak sayfaları içeren bir kopya verin.")
    print(f"[OK] {len(sayfalar)} sayfa · {len(metin):,} karakter · takma ad: {len(takma)}")
    llm.onay_al(f"'{rapor.name}' raporunun {len(sayfalar)} sayfası (uzun tablolarda ilk {ILK_SATIR} satır ve kodun özetleri) gönderilecek"
                f"{f'; {len(takma)} ad takma adla' if takma else ''}. Telefon, e-posta, TCKN, IBAN maskeli.", evet)
    izinli = izinli_sayilar(metin + " " + odak)
    adlar = {s["ad"] for s in sayfalar}
    n, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(metin, okuyucu, odak, d, o), lambda x: denetle(x, izinli, adlar), tur)
    n = geri_ac(n, takma)
    txt = rapor_yaz(cikti, n, kalan, yapilan, rapor.name, okuyucu)
    return {"not": n, "kalan": kalan, "tur": yapilan, "txt": txt, "sayfalar": sayfalar, "metin": metin, "takma": takma}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Excel raporundan yönetici için yorum, bulgu ve aksiyon notu yazar; nottaki rakamları raporla karşılaştırır.")
    p.add_argument("--rapor", type=Path, default=ORNEK / "satis_hedef_gerceklesme_raporu.xlsx", help="Excel rapor (.xlsx)")
    p.add_argument("--okuyucu", default="Satış Müdürü", help="Notun okuyucusu (ör. 'Şube Müdürü', 'Kredi Komitesi')")
    p.add_argument("--odak", default="", help="Notta öne çıkarılacak konu ya da soru")
    p.add_argument("--takma-ad", nargs="*", default=[], metavar="SÜTUN", help="Değerleri takma adla gönderilecek sütun başlıkları (ör. Temsilci Müşteri)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "rapor_notu.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.rapor, a.cikti, a.okuyucu, a.odak, a.takma_ad, a.tur, a.evet)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    n = s["not"]
    print(f"[OK] Not: {len(n['ozet'])} özet maddesi · {len(n['bulgular'])} bulgu · {len(n['aksiyonlar'])} aksiyon")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()} (+ {s['txt'].name})")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
