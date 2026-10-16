"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import json
import os
import sys
import tempfile
import unittest
from datetime import date, datetime
from decimal import Decimal as D
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
SATIS = ORNEK / "satis_hedef_gerceklesme_raporu.xlsx"
KREDI = ORNEK / "kredi_gecikme_yaslandirma_raporu.xlsx"
CEVAP = json.loads((Path(__file__).parent / "sahte_cevaplar.json").read_text(encoding="utf-8"))
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    """İlk not bilerek raporda olmayan bir fark (772.971) hesaplar ve takma ada ek getirir ('Kişi-4'ün')."""
    MESAJLAR.append(kullanici)
    return copy.deepcopy(CEVAP["dogru" if "<duzeltme>" in kullanici else "hatali"])


class KodTesti(unittest.TestCase):
    def test_bicim_ve_sayilar(self):
        self.assertEqual([agent.bicimle(*x) for x in ((0.5557638, "0.0%"), (-0.109289, "0%"), (967029.0800002, "#,##0"), (1740000, "General"),
                                                       (datetime(2026, 10, 15), ""), (date(2026, 1, 2), ""), (True, ""), ("  a  b ", ""), (None, ""))],
                         ["%55,6", "-%10,9", "967.029,08", "1.740.000", "15.10.2026", "02.01.2026", "Evet", "a b", ""])
        self.assertEqual(agent.sayi_tokenlari("Satış 967.029 TL, %55,6; tahmin 1.934.058,16 · 11/22 iş günü · 15.10.2026 · 2026-10 · 3.5"),
                         [(D(15), 0), (D(10), 0), (D(2026), 0), (D(967029), 0), (D("55.6"), 1), (D("1934058.16"), 2), (D(11), 0), (D(22), 0),
                          (D(2026), 0), (D(10), 0), (D("3.5"), 1)])
        izinli = {D("967029.08"), D("55.6"), D("-34.6")}
        self.assertTrue(all(agent.eslesir(*x, izinli) for x in ((D(967029), 0), (D("967029.1"), 1), (D(56), 0), (D("34.6"), 1))))
        self.assertFalse(any(agent.eslesir(*x, izinli) for x in ((D(967030), 0), (D("55.7"), 1), (D(772971), 0))))

    def test_okuma_ve_takma_ad(self):
        s, t = agent.sayfalari_oku(SATIS, ["Temsilci"])
        self.assertEqual([x["ad"] for x in s], ["Genel", "Temsilci", "Bölge", "Temsilci × Ürün Grubu", "Bilgi"])
        self.assertEqual(t, {"Ece Yıldız": "Kişi-1", "Selin Kaya": "Kişi-2", "Burak Şahin": "Kişi-3", "Ali Demir": "Kişi-4"})
        m = agent.rapor_metni(s)
        self.assertIn("Toplam | 967.029,08 | 1.740.000 | %55,6 | 883.425,94 | %9,5", m)
        self.assertIn("Kişi-4 | 181.654,17 | 490.000 | %37,1 | 277.724,53 | -%34,6", m)
        for ad in t:
            self.assertNotIn(ad, m)
        s, t = agent.sayfalari_oku(KREDI, ["Müşteri"])
        self.assertEqual(len(t), 17)                                  # Gecikmeli Krediler, Müşteri Bazında, Aralık Değiştirenler
        self.assertNotIn("Müşteri 20540", agent.rapor_metni(s))
        with tempfile.TemporaryDirectory() as tmp:                     # uzun tablo: ilk satırlar + kodun sütun özetleri
            y = Path(tmp) / "uzun.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Liste"
            ws.append(["Mağaza", "Satış", "Pay"])
            for i in range(1, 61):
                ws.append([f"M{i:02d}", i * 1000, i / 1830])
                ws.cell(ws.max_row, 3).number_format = "0.0%"
            wb.save(y)
            s, _ = agent.sayfalari_oku(y, [])
            self.assertEqual((len(s[0]["satirlar"]), s[0]["kesik"]), (1 + agent.ILK_SATIR, 60 - agent.ILK_SATIR))
            self.assertEqual(s[0]["ozet"][0], "Satış: toplam 1.830.000 · ortalama 30.500 · en büyük 60.000 (M60) · en küçük 1.000 (M01) · 60 değer")
            self.assertIn("en büyük %3,3 (M60)", s[0]["ozet"][1])
            self.assertIn("1.830.000", agent.rapor_metni(s))
        with self.assertRaises(llm.LLMHatasi):
            agent.sayfalari_oku(ORNEK / "yok.xlsx", [])

    def test_denetim(self):
        s, t = agent.sayfalari_oku(SATIS, ["Temsilci"])
        m = agent.rapor_metni(s)
        izinli, adlar = agent.izinli_sayilar(m), {x["ad"] for x in s}
        self.assertEqual(agent.denetle(CEVAP["dogru"], izinli, adlar), [])
        h = agent.denetle(CEVAP["hatali"], izinli, adlar)
        self.assertEqual(len(h), 4)
        self.assertTrue(h[0].startswith("Özet 3: raporda olmayan sayı: 772.971"))
        self.assertEqual(sum("takma ada ek getirilmiş" in x for x in h), 3)
        c = copy.deepcopy(CEVAP["dogru"])
        c["ozet"] = c["ozet"] * 2
        c["bulgular"][0]["sayfa"] = "Özet"
        c["aksiyonlar"][0]["bulgu_no"] = 9
        c["veri_notlari"].append("Telefon: [TELEFON]")
        h = agent.denetle(c, izinli, adlar)
        self.assertEqual(len(h), 4)
        self.assertTrue(any("8 madde" in x for x in h) and any("'Özet' raporda yok" in x for x in h) and any("bulgu_no 9" in x for x in h)
                        and any("maskeli yer tutucu" in x for x in h))

    def test_geri_ac(self):
        n = agent.geri_ac({"a": ["Kişi-1 ve Kişi-12"], "b": "Kişi-1:"}, {"Ece": "Kişi-1", "Can": "Kişi-12"})
        self.assertEqual(n, {"a": ["Ece ve Can"], "b": "Ece:"})


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ornek_ve_duzeltme(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "n.xlsx"
            s = agent.calistir(SATIS, cikti, "Satış Müdürü", "ay sonu hedef riski", ["Temsilci"], evet=True)
            self.assertEqual((len(MESAJLAR), s["tur"], s["kalan"]), (2, 1, []))
            self.assertIn("772.971", MESAJLAR[1].split("<duzeltme>")[1])
            self.assertIn("<odak>ay sonu hedef riski</odak>", MESAJLAR[0])
            for ad in ("Ali Demir", "Ece Yıldız"):
                self.assertNotIn(ad, MESAJLAR[0])
            txt = s["txt"].read_text(encoding="utf-8")
            self.assertIn("1. [Kritik] Ali Demir ekimde 181.654 TL", txt)
            self.assertNotIn("Kişi-", txt)
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Not", "Bulgular", "Aksiyonlar", "Kontrol"])
            self.assertEqual([wb["Bulgular"].cell(2, j).value for j in (2, 4)], ["Kritik", "Temsilci"])
            self.assertEqual(wb["Aksiyonlar"].max_row, 5)
            self.assertIn("sorunsuz", wb["Kontrol"].cell(2, 1).value)

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(SATIS, Path(tmp) / "n.xlsx", takma_ad=["Temsilci"], tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 4)
            self.assertIn("772.971", load_workbook(Path(tmp) / "n.xlsx")["Kontrol"].cell(2, 1).value)
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--rapor", str(Path(tmp) / "yok.xlsx"), "--evet"]), 1)
            (Path(tmp) / "r.csv").write_text("a;b\n1;2\n", encoding="utf-8")
            self.assertEqual(agent.main(["--rapor", str(Path(tmp) / "r.csv"), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
