"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import json
import os
import re
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
CEVAP = json.loads((Path(__file__).parent / "sahte_cevaplar.json").read_text(encoding="utf-8"))
SOZ = (ORNEK / "sozlesme_maddeleri.txt").read_text(encoding="utf-8")
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    """Yazıda ilk turda sözleşmede olmayan 'Madde 30.3'e atıf yapar."""
    MESAJLAR.append(kullanici)
    if "<ozet>" in kullanici:
        return copy.deepcopy(CEVAP["yazi" if "<duzeltme>" in kullanici else "yazi_hatali"])
    t = re.search(r'<rapor tarih="(\d\d)\.(\d\d)\.(\d{4})"', kullanici)
    return copy.deepcopy(CEVAP["olaylar"][f"{t[3]}-{t[2]}-{t[1]}.txt"])


def hesap(talep=None, basvuru=20):
    r = agent.raporlari_oku(ORNEK / "raporlar")
    o = {x["no"]: CEVAP["olaylar"][x["dosya"]] for x in r}
    return r, o, agent.gun_hesabi(r, o, talep or agent.VARSAYILAN_TALEP, basvuru)


class KodTesti(unittest.TestCase):
    def test_okuma(self):
        r = agent.raporlari_oku(ORNEK / "raporlar")
        self.assertEqual([(x["no"], x["tarih"]) for x in r][:2], [("R1", date(2026, 10, 5)), ("R2", date(2026, 10, 6))])
        with tempfile.TemporaryDirectory() as tmp:
            for ad in ("a.txt", "b.txt"):
                (Path(tmp) / ad).write_text("Tarih: 05.10.2026\nYağmur\n", encoding="utf-8")
            with self.assertRaises(llm.LLMHatasi):
                agent.raporlari_oku(Path(tmp))
            (Path(tmp) / "b.txt").write_text("tarihsiz\n", encoding="utf-8")
            with self.assertRaises(llm.LLMHatasi):
                agent.raporlari_oku(Path(tmp))

    def test_gun_hesabi(self):
        _, _, h = hesap()
        self.assertEqual((h["toplam"], h["tam"], h["kismi"], h["eszamanli"], h["haric"], h["basvuru_son"]),
                         (4, 2, 2, [date(2026, 10, 8)], [], date(2026, 10, 25)))
        self.assertEqual(h["neden"], {"Hava muhalefeti": {"Tam": 2, "Kısmi": 0}, "İdare kaynaklı": {"Tam": 0, "Kısmi": 2}})
        _, _, h = hesap(["Hava muhalefeti"])                               # idare talep dışı: 08.10 yalnız tedarik kaynaklı kalır
        self.assertEqual((h["toplam"], h["haric"]), (2, [date(2026, 10, 7), date(2026, 10, 8)]))
        o = copy.deepcopy(CEVAP["olaylar"])
        o["2026-10-06.txt"]["olaylar"][1]["durma"] = "Kısmi"                # aynı gün iki talep nedeni: gün bir kez sayılır
        r = agent.raporlari_oku(ORNEK / "raporlar")
        h = agent.gun_hesabi(r, {x["no"]: o[x["dosya"]] for x in r}, agent.VARSAYILAN_TALEP, None)
        self.assertEqual((h["toplam"], h["neden"]["İdare kaynaklı"]["Kısmi"], h["basvuru_son"]), (4, 3, None))
        self.assertIn("08.10.2026: Kısmi duruş · İdare kaynaklı", agent.ozet_metni(hesap()[2]))

    def test_denetim(self):
        r, o, h = hesap()
        self.assertEqual([agent.denetle_olay(x, o[x["no"]]) for x in r], [[]] * 5)
        c = copy.deepcopy(o["R1"])
        c["olaylar"][0]["alinti"] = "Saha tamamen sular altında kaldı."
        self.assertIn("raporda aynen geçmiyor", agent.denetle_olay(r[0], c)[0])
        self.assertEqual(agent.denetle_yazi(CEVAP["yazi"], h, r, SOZ), [])
        self.assertEqual(agent.denetle_yazi(CEVAP["yazi_hatali"], h, r, SOZ), ["Yazı: sözleşme metninde olmayan madde: 30.3"])
        y = copy.deepcopy(CEVAP["yazi"])
        y["govde"] = y["govde"].replace("toplam 4 gün", "toplam 6 gün").replace("09.10.2026 tarihli", "12.10.2026 tarihli") + " Aksi hâlde dava açılacaktır."
        s = agent.denetle_yazi(y, h, r, SOZ)
        self.assertTrue(any("olmayan sayı / tarih: 12.10.2026, 6" in x for x in s) and any("(4) gövdede yok" in x for x in s) and any("tehdit" in x for x in s))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ornek_ve_duzeltme(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "s.xlsx"
            self.assertEqual(agent.main(["--evet", "--cikti", str(cikti)]), 0)
            self.assertEqual(len(MESAJLAR), 7)                         # 5 rapor + yazı + yazı düzeltmesi
            self.assertIn("Talep edilebilir gün: 4 (tam duruş 2, kısmi duruş 2)", MESAJLAR[5])
            self.assertIn("Son başvuru tarihi (ilk olay + başvuru süresi): 25.10.2026", MESAJLAR[5])
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Olaylar", "Gün Hesabı", "Yazı", "Kontrol"])
            g = [list(x) for x in wb["Gün Hesabı"].iter_rows(min_row=2, max_row=5, values_only=True)]
            self.assertEqual([(x[1], x[2], x[4]) for x in g], [("Evet", "Tam", None), ("Evet", "Tam", None), ("Evet", "Kısmi", None), ("Evet", "Kısmi", "Evet")])
            self.assertIn("sorunsuz", wb["Kontrol"].cell(2, 1).value)
            y = "\n".join(str(c.value or "") for c in wb["Yazı"]["A"])
            self.assertIn("toplam 4 gün süre uzatımı", y)
            self.assertTrue(y.rstrip().endswith("Saygılarımızla,\nYüklenici Teknik Ofis"))

    @mock.patch.object(llm, "json_iste", side_effect=lambda *a, **k: {"olaylar": []})
    def test_gecikme_yok_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(ORNEK / "raporlar", Path(tmp) / "s.xlsx", ORNEK / "sozlesme_maddeleri.txt", evet=True)
            self.assertEqual((s["hesap"]["toplam"], s["yazi"]["govde"][:20]), (0, "Talep edilebilir gec"))
            self.assertEqual(agent.main(["--raporlar", str(Path(tmp) / "yok"), "--evet"]), 1)
            self.assertEqual(agent.main(["--sozlesme", str(Path(tmp) / "yok.txt"), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
