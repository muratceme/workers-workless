"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import re
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    idler = re.findall(r'<talep id="([^"]+)"', kullanici)
    # Model bilerek düşük öncelik veriyor; kod kuralları düzeltmeli
    cevap = []
    for i in idler:
        e, a, tur, ekip = "dusuk", "dusuk", "olay", "Hizmet Masası (1. seviye)"
        if i == "BT-3101":
            e, a = "dusuk", "yuksek"                         # tek kullanıcı, ay sonu → P3
        if i == "BT-3109":
            tur = "hizmet_talebi"
        cevap.append({"id": i, "kategori": "Diğer", "ekip": ekip, "etki": e, "aciliyet": a, "tur": tur, "ozet": "o", "ilk_yanit": "Merhaba",
                      "cozum_onerisi": "c"})
    return {"talepler": cevap}


class OnIslemeTesti(unittest.TestCase):
    def test_kesinti_ve_guvenlik(self):
        t = agent.talepleri_oku(ORNEK / "talepler.csv")
        olaylar = agent.on_isle(t, 30, 4)
        self.assertEqual([(o["id"], o["talepler"]) for o in olaylar], [("YK-1", ["BT-3102", "BT-3103", "BT-3104", "BT-3105", "BT-3106"])])
        g = {x["id"]: x for x in t}
        self.assertTrue(g["BT-3107"]["guvenlik"])
        self.assertEqual(g["BT-3108"]["konu_eslesme"], [])            # "kağıt" ağ sayılmaz

    def test_esik(self):
        t = agent.talepleri_oku(ORNEK / "talepler.csv")
        self.assertEqual(agent.on_isle(t, 30, 6), [])                  # 5 talep < 6

    def test_matris(self):
        self.assertEqual((agent.MATRIS[("yuksek", "yuksek")], agent.MATRIS[("orta", "orta")], agent.MATRIS[("dusuk", "dusuk")]), (1, 3, 5))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_kod_kurallari_ve_sla(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "d.xlsx"
            s = agent.calistir(ORNEK / "talepler.csv", cikti, ORNEK / "sla.json", evet=True)
            self.assertNotIn("ayse@ornek.com.tr", MESAJLAR[0])
            r = s["sonuc"]
            self.assertEqual(r["BT-3101"]["oncelik"], "P3")
            self.assertEqual(r["BT-3108"]["oncelik"], "P5")
            self.assertEqual(r["BT-3102"]["oncelik"], "P2")              # yaygın kesinti kuralı
            self.assertEqual((r["BT-3107"]["oncelik"], r["BT-3107"]["ekip"], r["BT-3107"]["tur"]), ("P2", "Bilgi Güvenliği", "guvenlik_olayi"))
            self.assertEqual(r["BT-3102"]["cozum_hedef"], datetime(2026, 11, 2, 17, 5))   # P2: 8 saat
            self.assertEqual(r["BT-3108"]["ilk_yanit_hedef"], datetime(2026, 11, 3, 11, 0))  # P5: 24 saat
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Yaygın Kesinti", "Talepler", "Özet"])
            self.assertEqual(wb["Talepler"].cell(2, 6).value, "P2")       # en yüksek öncelik en üstte


if __name__ == "__main__":
    unittest.main()
