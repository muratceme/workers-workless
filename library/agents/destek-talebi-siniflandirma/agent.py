"""
Destek Talebi Sınıflandırma — Workers / Workless AI Agent
Bilgi Teknolojileri › BT Destek Uzmanı

1. Kod talepleri okur, güvenlik olayı belirtilerini (oltalama, şifre girme, virüs, fidye yazılımı, şüpheli
   giriş) ve kısa sürede aynı konuda yoğunlaşan talepleri (olası yaygın kesinti: varsayılan 30 dakikada 4+
   benzer talep) bulur. E-posta, telefon, IBAN, TCKN maskelenir; gönderen modele gitmez.
2. Model her talebi kategori, ekip, etki, aciliyet ve türe (olay / hizmet talebi / güvenlik olayı) ayırır,
   kullanıcıya ilk yanıt taslağı ve destek uzmanına çözüm önerisi yazar.
3. Kod önceliği ITIL etki × aciliyet matrisiyle hesaplar (P1–P5), SLA dosyasından ilk yanıt ve çözüm hedef
   zamanlarını çıkarır. Kod kuralları modelin üstündedir: güvenlik olayları en az P2 ve Bilgi Güvenliği ekibi;
   yaygın kesinti kümesindeki talepler en az P2 ve tek ana olay altında gruplanır.

Kullanım:
    python agent.py                                           # örnek taleplerle dener
    python agent.py --girdi talepler.xlsx --sla sla.json
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import llm

BURASI = Path(__file__).resolve().parent
PAKET = 20
KATEGORILER = ["Şifre ve hesap kilidi", "Kullanıcı hesabı ve yetki", "Ağ ve internet", "E-posta ve iletişim", "ERP / iş uygulamaları",
               "Donanım (bilgisayar, monitör)", "Yazıcı ve çevre birimleri", "Yazılım kurulum ve lisans", "Güvenlik olayı", "Diğer"]
EKIPLER = ["Hizmet Masası (1. seviye)", "Sistem ve Ağ", "Uygulama Destek (ERP)", "Saha Destek (donanım)", "Bilgi Güvenliği"]
SEVIYE = ["dusuk", "orta", "yuksek"]
# ITIL örnek öncelik matrisi: (etki, aciliyet) → P
MATRIS = {("yuksek", "yuksek"): 1, ("yuksek", "orta"): 2, ("orta", "yuksek"): 2, ("yuksek", "dusuk"): 3, ("orta", "orta"): 3,
          ("dusuk", "yuksek"): 3, ("orta", "dusuk"): 4, ("dusuk", "orta"): 4, ("dusuk", "dusuk"): 5}
VARSAYILAN_SLA = {"P1": {"ilk_yanit": 0.25, "cozum": 4}, "P2": {"ilk_yanit": 1, "cozum": 8}, "P3": {"ilk_yanit": 4, "cozum": 24},
                  "P4": {"ilk_yanit": 8, "cozum": 72}, "P5": {"ilk_yanit": 24, "cozum": 168}}
GUVENLIK = re.compile(r"(oltalama|phishing|linke tıkla|bağlantıya tıkla|şifremi girdim|şifremi yazdım|virüs|fidye|ransomware|"
                      r"şüpheli giriş|hesabım ele geçir|dosyalarım şifrelen|tanımadığım bir yerden giriş)", re.I)
KESINTI_KONULARI = {
    # Kelime başı/sonu sınırlı: "ağ" kelimesi "kağıt", "Ağustos" içinde eşleşmesin
    "Ağ ve internet": re.compile(r"(\binternet|\bwi-?fi\b|\bağ(?:a|da|dan|ı)?\b|\bbağlantı|\bvpn\b)", re.I),
    "E-posta ve iletişim": re.compile(r"(\be-?posta|\boutlook\b|\be-?mail\b|\bteams\b)", re.I),
    "ERP / iş uygulamaları": re.compile(r"(\berp\b|\blogo\b|\bsap\b|\bnetsis\b|\bmikro\b|uygulama açılmıyor)", re.I),
}


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def zaman(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y"):
        try:
            return datetime.strptime(str(x).strip(), f)
        except ValueError:
            pass
    return None


def talepleri_oku(yol: Path) -> list[dict]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    s = [r for r in s if any(c not in (None, "") for c in r)]
    b = [kucuk(x) for x in s[0]]
    bul = lambda *a: next((i for i, x in enumerate(b) if x in a), None)  # noqa: E731
    k = {"id": bul("talep no", "kayıt no", "ticket", "id", "ıd", "no"), "tarih": bul("tarih", "oluşturma tarihi", "açılış"),
         "gonderen": bul("gönderen", "kullanıcı", "talep eden"), "departman": bul("departman", "birim", "lokasyon"),
         "konu": bul("konu", "başlık"), "mesaj": bul("mesaj", "açıklama", "içerik")}
    if k["mesaj"] is None or k["tarih"] is None:
        raise SystemExit(f"Tarih ve Mesaj sütunları gerekli. Başlıklar: {s[0]}")
    al = lambda r, a: r[k[a]] if k[a] is not None and k[a] < len(r) else None  # noqa: E731
    return [{"id": str(al(r, "id") or f"T{i}"), "zaman": zaman(al(r, "tarih")), "gonderen": str(al(r, "gonderen") or ""),
             "departman": str(al(r, "departman") or ""), "konu": str(al(r, "konu") or ""), "mesaj": str(al(r, "mesaj") or "").strip()}
            for i, r in enumerate(s[1:], 1) if str(al(r, "mesaj") or "").strip()]


def on_isle(talepler: list[dict], pencere_dk: int, esik: int) -> list[dict]:
    """Güvenlik belirtisi ve yaygın kesinti kümeleri. Dönüş: kesinti olayları."""
    for t in talepler:
        metin = f"{t['konu']} {t['mesaj']}"
        t["guvenlik"] = bool(GUVENLIK.search(metin))
        t["konu_eslesme"] = [k for k, d in KESINTI_KONULARI.items() if d.search(metin)]
        t["kesinti"] = None
    olaylar = []
    for konu in KESINTI_KONULARI:
        liste = sorted((t for t in talepler if konu in t["konu_eslesme"] and t["zaman"]), key=lambda t: t["zaman"])
        i = 0
        while i < len(liste):
            j = i
            while j + 1 < len(liste) and liste[j + 1]["zaman"] - liste[i]["zaman"] <= timedelta(minutes=pencere_dk):
                j += 1
            if j - i + 1 >= esik:
                grup = liste[i:j + 1]
                olay_id = f"YK-{len(olaylar) + 1}"
                for t in grup:
                    t["kesinti"] = olay_id
                olaylar.append({"id": olay_id, "konu": konu, "talepler": [t["id"] for t in grup], "bas": grup[0]["zaman"], "bit": grup[-1]["zaman"],
                                "departmanlar": sorted({t["departman"] for t in grup if t["departman"]})})
                i = j + 1
            else:
                i += 1
    return olaylar


SEMA = {
    "type": "object",
    "properties": {"talepler": {"type": "array", "items": {
        "type": "object",
        "properties": {
            "id": {"type": "string"},
            "kategori": {"type": "string", "enum": KATEGORILER},
            "ekip": {"type": "string", "enum": EKIPLER},
            "etki": {"type": "string", "enum": SEVIYE},
            "aciliyet": {"type": "string", "enum": SEVIYE},
            "tur": {"type": "string", "enum": ["olay", "hizmet_talebi", "guvenlik_olayi"]},
            "ozet": {"type": "string"},
            "ilk_yanit": {"type": "string"},
            "cozum_onerisi": {"type": "string"},
        },
        "required": ["id", "kategori", "ekip", "etki", "aciliyet", "tur", "ozet", "ilk_yanit", "cozum_onerisi"],
        "additionalProperties": False}}},
    "required": ["talepler"],
    "additionalProperties": False,
}


def siniflandir(talepler: list[dict]) -> dict[str, dict]:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    sonuc = {}
    for i in range(0, len(talepler), PAKET):
        parca = talepler[i:i + PAKET]
        satirlar = []
        for t in parca:
            ipucu = []
            if t["guvenlik"]:
                ipucu.append("güvenlik olayı belirtisi")
            if t["kesinti"]:
                ipucu.append(f"olası yaygın kesinti kümesi {t['kesinti']}")
            satirlar.append(f'<talep id="{t["id"]}" departman="{t["departman"]}" zaman="{t["zaman"]:%d.%m.%Y %H:%M}" ipuclari="{"; ".join(ipucu) or "-"}">\n'
                            f"{llm.maskele(t['konu'])}\n{llm.maskele(t['mesaj'])}\n</talep>")
        mesaj = "\n".join(["<kategoriler>" + "; ".join(KATEGORILER) + "</kategoriler>", "<ekipler>" + "; ".join(EKIPLER) + "</ekipler>",
                           "<talepler>", *satirlar, "</talepler>"])
        yanit = llm.json_iste(sistem, mesaj, SEMA)
        gecerli = {t["id"] for t in parca}
        for s in yanit.get("talepler", []):
            if s.get("id") in gecerli:
                sonuc[s["id"]] = s
    return sonuc


def oncelik_hesapla(t: dict, s: dict, sla: dict) -> list[str]:
    duzeltme = []
    if t["guvenlik"] or s["tur"] == "guvenlik_olayi":
        if s["ekip"] != "Bilgi Güvenliği":
            s["ekip"] = "Bilgi Güvenliği"
            duzeltme.append("ekip → Bilgi Güvenliği (kural)")
        if s["tur"] != "guvenlik_olayi":
            s["tur"] = "guvenlik_olayi"
            duzeltme.append("tür → güvenlik olayı (kural)")
    p = MATRIS[(s["etki"], s["aciliyet"])]
    if (s["tur"] == "guvenlik_olayi" or t["kesinti"]) and p > 2:
        duzeltme.append(f"öncelik P{p} → P2 ({'güvenlik olayı' if s['tur'] == 'guvenlik_olayi' else 'yaygın kesinti'} kuralı)")
        p = 2
    s["oncelik"] = f"P{p}"
    h = sla.get(s["oncelik"], VARSAYILAN_SLA[s["oncelik"]])
    s["ilk_yanit_hedef"] = t["zaman"] + timedelta(hours=h["ilk_yanit"]) if t["zaman"] else None
    s["cozum_hedef"] = t["zaman"] + timedelta(hours=h["cozum"]) if t["zaman"] else None
    return duzeltme


BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
ONAY_DOLGU = PatternFill("solid", fgColor="FFF4CE")
P_DOLGU = {"P1": "F8C4C1", "P2": "FDE2E1", "P3": "FFF4CE", "P4": "E3F5E1", "P5": "EDEDED"}
UST = Alignment(vertical="top", wrap_text=True)
TUR = {"olay": "Olay", "hizmet_talebi": "Hizmet talebi", "guvenlik_olayi": "GÜVENLİK OLAYI"}


def rapor_yaz(cikti: Path, talepler, sonuc, olaylar, duzeltmeler) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Talepler"
    bas = ["Talep No", "Zaman", "Departman", "Konu", "Mesaj", "Öncelik", "Tür", "Kategori", "Ekip", "Etki", "Aciliyet", "Ana Olay", "Özet",
           "İlk Yanıt Taslağı", "Çözüm Önerisi", "İlk Yanıt Hedefi", "Çözüm Hedefi", "Kural Düzeltmesi", "Onay"]
    ws.append(bas)
    for c in ws[1]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI
    for j, w in enumerate((10, 16, 12, 18, 50, 8, 16, 22, 22, 8, 9, 9, 40, 60, 50, 16, 16, 34, 8), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for t in sorted(talepler, key=lambda t: (sonuc.get(t["id"], {}).get("oncelik", "P9"), t["zaman"] or datetime.min)):
        s = sonuc.get(t["id"])
        if not s:
            ws.append([t["id"], t["zaman"], t["departman"], t["konu"], t["mesaj"], "", "", "", "", "", "", t["kesinti"] or "", "AI yanıt vermedi"])
            continue
        ws.append([t["id"], t["zaman"], t["departman"], t["konu"], t["mesaj"], s["oncelik"], TUR[s["tur"]], s["kategori"], s["ekip"], s["etki"],
                   s["aciliyet"], t["kesinti"] or "", s["ozet"], s["ilk_yanit"], s["cozum_onerisi"], s["ilk_yanit_hedef"], s["cozum_hedef"],
                   "; ".join(duzeltmeler.get(t["id"], [])), ""])
        r = ws.max_row
        for c in (2, 16, 17):
            ws.cell(r, c).number_format = "DD.MM.YYYY HH:MM"
        ws.cell(r, 6).fill = PatternFill("solid", fgColor=P_DOLGU[s["oncelik"]])
        for c in (8, 9, 10, 11, 13, 14, 15):
            ws.cell(r, c).fill = AI_DOLGU
        ws.cell(r, 19).fill = ONAY_DOLGU
        for c in ws[r]:
            c.alignment = UST
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    o = wb.create_sheet("Yaygın Kesinti", 0)
    o.append(["Ana Olay", "Konu", "Talep Sayısı", "İlk Talep", "Son Talep", "Etkilenen Departmanlar", "Talepler", "Yapılacak"])
    for c in o[1]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI
    for x in olaylar:
        o.append([x["id"], x["konu"], len(x["talepler"]), x["bas"], x["bit"], ", ".join(x["departmanlar"]), ", ".join(x["talepler"]),
                  "Ana olay kaydı açın, kullanıcılara toplu bilgilendirme yapın, tekil talepleri ana olaya bağlayın"])
        o.cell(o.max_row, 4).number_format = o.cell(o.max_row, 5).number_format = "DD.MM.YYYY HH:MM"
    if not olaylar:
        o.append(["—", "Yaygın kesinti belirtisi yok"])
    for j, w in enumerate((10, 22, 12, 16, 16, 30, 40, 70), 1):
        o.column_dimensions[get_column_letter(j)].width = w

    p = wb.create_sheet("Özet")
    p.append(["Öncelik", "Adet"])
    for c in p[1]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI
    for k, n in sorted(Counter(s["oncelik"] for s in sonuc.values()).items()):
        p.append([k, n])
    p.append([])
    p.append(["Ekip", "Adet"])
    for c in p[p.max_row]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI
    for k, n in Counter(s["ekip"] for s in sonuc.values()).most_common():
        p.append([k, n])
    p.append([])
    p.append(["Öncelik matrisi", "ITIL etki × aciliyet: Y×Y=P1; Y×O, O×Y=P2; Y×D, O×O, D×Y=P3; O×D, D×O=P4; D×D=P5"])
    p.append(["Kod kuralları", "Güvenlik olayı → Bilgi Güvenliği ve en az P2; yaygın kesinti kümesi → en az P2"])
    p.append(["Model", llm.kullanim_ozeti()])
    p.column_dimensions["A"].width = 26
    p.column_dimensions["B"].width = 100
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(girdi: Path, cikti: Path, sla_yolu: Path | None = None, pencere_dk: int = 30, esik: int = 4, evet: bool = False) -> dict:
    talepler = talepleri_oku(girdi)
    if not talepler:
        raise llm.LLMHatasi(f"{girdi.name}: talep yok.")
    sla = dict(VARSAYILAN_SLA)
    if sla_yolu:
        sla.update({k: v for k, v in json.loads(sla_yolu.read_text(encoding="utf-8")).items() if k.startswith("P")})
    olaylar = on_isle(talepler, pencere_dk, esik)
    print(f"[OK] {len(talepler)} talep · güvenlik belirtisi: {sum(t['guvenlik'] for t in talepler)} · olası yaygın kesinti: {len(olaylar)}")
    llm.onay_al(f"{len(talepler)} talebin konu ve mesajı (e-posta/telefon/IBAN/TCKN maskeli; gönderen hariç) gönderilecek.", evet)
    sonuc = siniflandir(talepler)
    duzeltmeler = {t["id"]: oncelik_hesapla(t, sonuc[t["id"]], sla) for t in talepler if t["id"] in sonuc}
    rapor_yaz(cikti, talepler, sonuc, olaylar, duzeltmeler)
    return {"talepler": talepler, "sonuc": sonuc, "olaylar": olaylar, "duzeltmeler": duzeltmeler}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="BT destek taleplerini sınıflandırır, ITIL önceliği ve SLA hedeflerini çıkarır.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "talepler.csv",
                   help="Talepler (.xlsx/.csv): Tarih, Mesaj; Talep No, Gönderen, Departman, Konu isteğe bağlı")
    p.add_argument("--sla", type=Path, help="Öncelik bazında SLA saatleri (JSON)")
    p.add_argument("--pencere", type=int, default=30, help="Yaygın kesinti penceresi, dakika (varsayılan 30)")
    p.add_argument("--esik", type=int, default=4, help="Pencerede yaygın kesinti sayılacak en az talep (varsayılan 4)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "destek_talepleri.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    if a.girdi == BURASI / "ornek_veri" / "talepler.csv":
        a.sla = a.sla or BURASI / "ornek_veri" / "sla.json"
    try:
        s = calistir(a.girdi, a.cikti, a.sla, a.pencere, a.esik, a.evet)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    say = Counter(x["oncelik"] for x in s["sonuc"].values())
    print(f"[OK] {len(s['sonuc'])}/{len(s['talepler'])} talep · " + " ".join(f"{k}:{n}" for k, n in sorted(say.items())))
    for o in s["olaylar"]:
        print(f"[!] {o['id']} {o['konu']}: {len(o['talepler'])} talep {o['bas']:%H:%M}–{o['bit']:%H:%M}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
