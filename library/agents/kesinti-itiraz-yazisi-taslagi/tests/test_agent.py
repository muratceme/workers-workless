"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import json
import os
import re
import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
BUGUN = date(2026, 10, 16)
CEVAP = json.loads((Path(__file__).parent / "sahte_cevaplar.json").read_text(encoding="utf-8"))
DAYANAK = (ORNEK / "dayanaklar.txt").read_text(encoding="utf-8")
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    """F-202 için ilk turda dayanakta olmayan 'Madde 4.5'e atıf yapar."""
    MESAJLAR.append(kullanici)
    f = re.search(r'fatura="([^"]+)"', kullanici)[1]
    if f == "F-202" and "<duzeltme>" not in kullanici:
        return copy.deepcopy(CEVAP["F-202_hatali"])
    return copy.deepcopy(CEVAP[f])


def kesintiler():
    _, k = agent.kesintileri_oku(ORNEK / "kesintiler.csv")
    for x in k:
        x["belgeler"] = agent.belgeleri_oku(ORNEK / "belgeler", x["fatura"], {})
    return {x["fatura"]: x for x in k}


class KodTesti(unittest.TestCase):
    def test_okuma(self):
        b, k = agent.kesintileri_oku(ORNEK / "kesintiler.csv")
        self.assertEqual(b[:5], ["Sigorta Şirketi", "Fatura No", "Kesinti Tarihi", "Kesinti Tutarı", "Kesinti Nedeni"])
        self.assertEqual([(x["fatura"], x["tutar"], x["tarih"]) for x in k],
                         [("F-201", D("4000.00"), date(2026, 10, 5)), ("F-202", D("2500.00"), date(2026, 10, 7)), ("F-203", D("1800.00"), date(2026, 10, 9))])
        k = kesintiler()
        self.assertEqual((list(k["F-201"]["belgeler"]), list(k["F-203"]["belgeler"])), (["F-201.txt"], []))
        self.assertNotIn("provizyon@sigortab", k["F-202"]["belgeler"]["F-202.txt"])          # e-posta maskeli
        self.assertEqual(agent.sureleri_oku(ORNEK / "itiraz_sureleri.csv"), {"sigorta a": 30, "sigorta b": 15})
        self.assertEqual(agent.maddeler("Madde 7.2 ve md. 4.3, SUT 2.4.4.D değil, SUT 2.4.4 ve Ek-2"), {"7.2", "4.3", "2.4.4", "ek-2"})

    def test_sure(self):
        k = kesintiler()
        s = agent.sureleri_oku(ORNEK / "itiraz_sureleri.csv")
        self.assertEqual(agent.sure_durumu(k["F-202"], s, BUGUN), (date(2026, 10, 22), "Son itiraz tarihi 22.10.2026 (6 gün)"))
        self.assertEqual(agent.sure_durumu(k["F-202"], s, date(2026, 10, 18))[1], "İtiraz süresinin bitmesine 4 gün var (22.10.2026)")
        self.assertEqual(agent.sure_durumu(k["F-202"], s, date(2026, 10, 23))[1], "İtiraz süresi geçmiş (22.10.2026)")
        self.assertEqual(agent.sure_durumu(k["F-202"], {}, BUGUN), (None, ""))

    def test_denetim(self):
        k = kesintiler()
        self.assertEqual([agent.denetle(k[f], CEVAP[f], DAYANAK) for f in ("F-201", "F-202", "F-203")], [[], [], []])
        self.assertEqual(agent.denetle(k["F-202"], CEVAP["F-202_hatali"], DAYANAK),
                         ["K2: dayanak metninde ya da belgelerde olmayan madde / ek: 4.5 (yalnız verilen maddelere dayanın)"])
        c = copy.deepcopy(CEVAP["F-201"])
        c["itiraz_metni"] = c["itiraz_metni"].replace("4.000,00 TL kesinti", "4.500,00 TL kesinti").replace("F-201 numaralı", "Bu") + " Aksi hâlde dava açılacaktır."
        c["kullanilan_belgeler"] = ["F-201.txt", "sozlesme.pdf"]
        s = agent.denetle(k["F-201"], c, DAYANAK)
        for parca in ("olmayan sayı / tarih: 4500", "sozlesme.pdf", "fatura no 'F-201'", "'dava açılacaktır'"):
            self.assertTrue(any(parca in x for x in s), parca)
        c = copy.deepcopy(CEVAP["F-203"]) | {"oneri": "Evet", "itiraz_metni": "F-203 için 1.800,00 TL kesintinin kaldırılmasını arz ederiz."}
        self.assertIn("belgeye dayanmalı", agent.denetle(k["F-203"], c, DAYANAK)[0])
        c = copy.deepcopy(CEVAP["F-203"]) | {"eksik_belgeler": []}
        self.assertIn("eksik belgeler yazılmalı", agent.denetle(k["F-203"], c, DAYANAK)[0])


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ornek_ve_duzeltme(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "i.xlsx"
            self.assertEqual(agent.main(["--evet", "--cikti", str(cikti)]), 0)
            self.assertEqual(len(MESAJLAR), 4)                         # 3 kesinti + F-202 düzeltmesi
            self.assertIn("Madde 4.5", MESAJLAR[3].split("<duzeltme>")[1])
            self.assertIn("(bu fatura için belge verilmedi)", MESAJLAR[2])
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Kesintiler", "İtiraz Yazıları", "Kontrol"])
            ws = wb["Kesintiler"]
            self.assertEqual([c.value for c in ws[1]][:9], ["Sigorta Şirketi", "Fatura No", "Kesinti Tarihi", "Kesinti Tutarı", "Kesinti Nedeni", "İşlem",
                                                           "İtiraz Tarihi", "İtiraz Sonucu", "Geri Alınan Tutar"])
            self.assertEqual([ws.cell(r, 10).value for r in range(2, 5)], ["Evet", "Evet", "Belge eksik"])
            self.assertEqual(ws.cell(3, 16).value, "Son itiraz tarihi 22.10.2026 (6 gün)")
            y = wb["İtiraz Yazıları"]
            self.assertEqual(y.max_row, 3)
            self.assertTrue(y.cell(3, 5).value.endswith("Saygılarımızla,\nAnlaşmalı Kurumlar Birimi"))
            self.assertIn("Madde 4.3", y.cell(3, 5).value)

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_gizle_ve_cli(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            b = Path(tmp) / "belgeler"
            b.mkdir()
            (b / "F-201_epikriz.txt").write_text("Hasta Ayşe Deneme, kolesistektomi. Ek-2 fiyat listesi 38.500,00 TL.\n", encoding="utf-8")
            s = agent.calistir(ORNEK / "kesintiler.csv", Path(tmp) / "i.xlsx", b, ORNEK / "dayanaklar.txt", None, BUGUN, ["Ayşe Deneme"], evet=True)
            self.assertNotIn("Ayşe", "".join(MESAJLAR))
            self.assertIn("[HASTA-1]", MESAJLAR[0])
            self.assertEqual(s["kesintiler"][0]["sure"], "")
            self.assertEqual(agent.main(["--tarih", "x", "--evet"]), 2)
            self.assertEqual(agent.main(["--kesintiler", str(Path(tmp) / "yok.csv"), "--evet"]), 1)
            self.assertEqual(agent.main(["--dayanak", str(Path(tmp) / "yok.txt"), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
