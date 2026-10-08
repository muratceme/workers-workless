"""
İhale Dokümanı Özeti — Workers / Workless AI Agent
İnşaat › İhale ve Teklif › İhale Mühendisi

1. Kod ihale dokümanını (idari şartname, sözleşme tasarısı, teknik şartname, ilan, zeyilname; .pdf/.docx/.txt)
   okur ve şunları kendisi tarar: işin süresi ve gecikme cezası oranının belgeler arasında tutarlı olup
   olmadığı, ihale tarihi, belgelerde geçen önemli konular (teminat, iş deneyimi, fiyat farkı, avans, sigorta...).
2. Model; ihale bilgilerini, yeterlik kriterlerini, teminatları, süreleri, cezaları, istenen belgeleri ve özel
   şartları/riskleri **kaynak madde ve kelimesi kelimesine alıntıyla** özetler; açıklama talebi soruları önerir.
3. Kod her alıntıyı belgede arar (bulunamayan işaretlenir), oranların alıntıda geçtiğini kontrol eder, modelin
   atladığı konuları listeler ve --teklif verilirse tutarları hesaplar: geçici/kesin teminat, iş deneyimi ve iş
   hacmi alt sınırları, günlük gecikme cezası.

Kullanım:
    python agent.py                                              # örnek: kurgusal pazar yeri yapım ihalesi
    python agent.py --girdi ./ihale_dokumani --teklif 48500000 --bugun 08.10.2026
"""
from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
PAKET_KARAKTER = 120000
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")

BILGI_ALANLARI = ["İdare", "İhale kayıt numarası", "İşin adı", "İşin yeri", "İşin miktarı / kapsamı", "İhale usulü", "Teklif türü",
                  "İhale tarihi ve saati", "Teklif verme yeri / yöntemi", "İşin süresi", "Teklif geçerlilik süresi", "Yaklaşık maliyet",
                  "Kısmi teklif", "İş ortaklığı / konsorsiyum", "Alt yüklenici", "Fiyat farkı", "Avans", "Ödeme / hakediş",
                  "Yerli istekli fiyat avantajı", "Aşırı düşük teklif / sınır değer", "Değerlendirme yöntemi", "Benzer iş tanımı"]
KRITER_GRUPLARI = ["Belge / mesleki", "İş deneyimi", "Ekonomik ve mali", "Teknik (personel, makine, kalite)", "Diğer"]
TEMINAT_TURLERI = ["Geçici teminat", "Kesin teminat", "Ek kesin teminat", "Avans teminatı", "Diğer"]
TABANLAR = ["teklif bedeli", "sözleşme bedeli", "yaklaşık maliyet", "diğer", "yok"]
BIRIMLER = {"yüzde": Decimal(100), "binde": Decimal(1000), "on binde": Decimal(10000)}
# Kodun belgede aradığı konular: (konu, anahtar ifadeler — katlanmış metinde aranır)
KONULAR = [
    ("Geçici teminat", ("gecici teminat",)), ("Kesin teminat", ("kesin teminat",)), ("İş deneyimi", ("is deneyim",)),
    ("Bilanço / iş hacmi", ("bilanco", "is hacmi", "ciro")), ("Anahtar teknik personel", ("teknik personel",)),
    ("Gecikme cezası", ("gecikme ceza",)), ("Fiyat farkı", ("fiyat fark",)), ("Avans", ("avans",)), ("Alt yüklenici", ("alt yuklenici",)),
    ("Teklif geçerlilik süresi", ("gecerlilik sure",)), ("Yer teslimi", ("yer teslim",)), ("İş artışı", ("is artis",)),
    ("Sigorta", ("sigorta",)), ("Geçici / kesin kabul", ("gecici kabul", "kesin kabul")), ("Hakediş / ödeme", ("hakedis", "odeme")),
    ("Açıklama talebi", ("aciklama talep",)), ("Kalite belgesi", ("iso 9001", "kalite yonetim")),
]


def katla(s) -> str:
    return " ".join(str(s or "").translate(_TR).lower().split())


def duz(s) -> str:
    """Alıntı karşılaştırması için: Türkçe karakter katlanmış, noktalama atılmış, tek boşluklu metin."""
    return " ".join(re.sub(r"[^0-9a-z%]+", " ", katla(s)).split())


def tr_sayi(s: str) -> Decimal | None:
    s = s.strip()
    if re.fullmatch(r"\d{1,3}(\.\d{3})+(,\d+)?", s):
        s = s.replace(".", "")
    s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def sayilar(metin: str) -> set[Decimal]:
    return {d for d in (tr_sayi(x) for x in re.findall(r"\d+(?:[.,]\d+)*", metin or "")) if d is not None}


def tl(x: Decimal) -> str:
    return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".") + " TL"


# ----------------------------------------------------------------------------
# Belgeler ve kod taramaları
# ----------------------------------------------------------------------------

@dataclass
class Dosya:
    ad: str
    tur: str
    metin: str


def belge_turu(ad: str, metin: str) -> str:
    a, bas = katla(ad), katla(metin[:400])
    for anahtar, tur in (("zeyil", "Zeyilname"), ("idari", "İdari şartname"), ("sozlesme", "Sözleşme tasarısı"), ("teknik", "Teknik şartname"),
                         ("ilan", "İhale ilanı"), ("birim fiyat", "Birim fiyat / mahal listesi"), ("mahal", "Birim fiyat / mahal listesi")):
        if anahtar in a or anahtar.replace(" ", "_") in a:
            return tur
    for anahtar, tur in (("zeyilname", "Zeyilname"), ("idari sartname", "İdari şartname"), ("sozlesme tasarisi", "Sözleşme tasarısı"),
                         ("teknik sartname", "Teknik şartname"), ("ilan", "İhale ilanı")):
        if anahtar in bas:
            return tur
    return "Diğer"


def dosyalari_oku(girdi: Path) -> tuple[list[Dosya], list[str]]:
    yollar = sorted(girdi.iterdir()) if girdi.is_dir() else [girdi]
    dosyalar, atlanan = [], []
    for yol in yollar:
        if yol.is_dir() or yol.name.startswith((".", "~$")):
            continue
        try:
            metin = belge.metin_oku(yol)
        except belge.BelgeHatasi as h:
            atlanan.append(f"{yol.name} ({h})")
            continue
        if metin.strip():
            dosyalar.append(Dosya(yol.name, belge_turu(yol.name, metin), metin))
    return dosyalar, atlanan


def cumleler(metin: str) -> list[str]:
    return [c.strip() for c in re.split(r"(?<=[.;:])\s+(?=[A-ZÇĞİÖŞÜ0-9])|\n", metin) if c.strip()]


def tutarlilik(dosyalar: list[Dosya]) -> list[tuple[str, str, str]]:
    """İşin süresi ve gecikme cezası oranı belgeler arasında farklı mı? (önem, konu, açıklama)"""
    bulgular = []
    sure, ceza = {}, {}
    for d in dosyalar:
        for c in cumleler(d.metin):
            k = katla(c)
            if ("isin suresi" in k or "isi bitirme" in k) and "takvim gun" in k:
                for m in re.finditer(r"(\d+)\s*(?:\([^)]*\)\s*)?takvim gun", k):
                    sure.setdefault(int(m.group(1)), []).append(d.ad)
            if "gecikme" in k and "ceza" in k:
                for m in re.finditer(r"(on binde|binde|yuzde|%)\s*(\d+(?:[.,]\d+)?)", k):
                    birim = {"%": "yüzde", "yuzde": "yüzde", "binde": "binde", "on binde": "on binde"}[m.group(1)]
                    if re.search(r"\bher\b.*\bgun", k):                       # yalnız günlük oranlar
                        ceza.setdefault(f"{birim} {m.group(2)}", []).append(d.ad)
    if len(sure) > 1:
        bulgular.append(("Hata", "İşin süresi", "Belgeler arasında farklı süre: " +
                         "; ".join(f"{g} takvim günü ({', '.join(sorted(set(a)))})" for g, a in sorted(sure.items())) +
                         ". Açıklama talebiyle netleştirin; genellikle idari şartname ile sözleşme tasarısı aynı olmalıdır."))
    if len(ceza) > 1:
        bulgular.append(("Hata", "Gecikme cezası", "Belgeler arasında farklı günlük ceza oranı: " +
                         "; ".join(f"{o} ({', '.join(sorted(set(a)))})" for o, a in ceza.items())))
    return bulgular


def konu_taramasi(dosyalar: list[Dosya]) -> dict[str, list[str]]:
    sonuc = {}
    for konu, anahtarlar in KONULAR:
        for d in dosyalar:
            eslesen = [c for c in cumleler(d.metin) if any(a in katla(c) for a in anahtarlar)]
            if eslesen:                                        # madde başlığı yerine ilk gerçek cümleyi göster
                c = next((x for x in eslesen if len(x) > 40), eslesen[0])
                sonuc.setdefault(konu, []).append(f"{d.ad}: {c[:90]}{'…' if len(c) > 90 else ''}")
    return sonuc


def ihale_tarihi_bul(dosyalar: list[Dosya]) -> datetime | None:
    for d in dosyalar:
        m = re.search(r"ihale tarihi\s*(?:ve saati)?\s*:?\s*(\d{2}\.\d{2}\.\d{4})\D{0,12}(\d{2})[:.](\d{2})", katla(d.metin))
        if m:
            try:
                return datetime.strptime(f"{m.group(1)} {m.group(2)}:{m.group(3)}", "%d.%m.%Y %H:%M")
            except ValueError:
                pass
    return None


def gizle(metin: str, terimler: list[str]) -> str:
    metin = llm.maskele(metin)
    for i, t in enumerate(sorted({t.strip() for t in terimler if t.strip()}, key=len, reverse=True), 1):
        metin = re.sub(re.escape(t), f"[GİZLİ-{i}]", metin, flags=re.I)
    return metin


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

def _kalem(ozellikler: dict) -> dict:
    ozellikler = {**ozellikler, "kaynak": {"type": "string"}, "alinti": {"type": "string"}}
    return {"type": "array", "items": {"type": "object", "properties": ozellikler, "required": list(ozellikler), "additionalProperties": False}}


S = {"type": "string"}
ORAN = {"type": ["number", "null"]}
SEMA = {
    "type": "object",
    "properties": {
        "ihale_bilgileri": _kalem({"alan": {"type": "string", "enum": BILGI_ALANLARI}, "deger": S}),
        "yeterlik_kriterleri": _kalem({"grup": {"type": "string", "enum": KRITER_GRUPLARI}, "kriter": S, "aciklama": S, "oran_yuzde": ORAN,
                                       "oran_tabani": {"type": "string", "enum": TABANLAR}}),
        "teminatlar": _kalem({"tur": {"type": "string", "enum": TEMINAT_TURLERI}, "aciklama": S, "oran_yuzde": ORAN,
                              "oran_tabani": {"type": "string", "enum": TABANLAR}}),
        "sureler": _kalem({"olay": S, "deger": S}),
        "cezalar": _kalem({"tur": S, "aciklama": S, "oran": ORAN, "birim": {"type": "string", "enum": [*BIRIMLER, "TL", "diğer"]},
                           "periyot": {"type": "string", "enum": ["günlük", "bir kez", "diğer"]},
                           "oran_tabani": {"type": "string", "enum": TABANLAR}}),
        "istenen_belgeler": _kalem({"belge": S}),
        "ozel_sartlar": _kalem({"konu": S, "aciklama": S, "risk": {"type": "string", "enum": ["dusuk", "orta", "yuksek"]}, "neden": S}),
        "aciklama_talebi_sorulari": {"type": "array", "items": S},
        "genel_ozet": S,
    },
    "required": ["ihale_bilgileri", "yeterlik_kriterleri", "teminatlar", "sureler", "cezalar", "istenen_belgeler", "ozel_sartlar",
                 "aciklama_talebi_sorulari", "genel_ozet"],
    "additionalProperties": False,
}
LISTELER = ["ihale_bilgileri", "yeterlik_kriterleri", "teminatlar", "sureler", "cezalar", "istenen_belgeler", "ozel_sartlar"]


def paketle(dosyalar: list[Dosya], terimler: list[str]) -> list[list[tuple[Dosya, str]]]:
    paketler, simdiki, boy = [], [], 0
    for d in dosyalar:
        metin = gizle(d.metin, terimler)
        parcalar = [metin[i:i + PAKET_KARAKTER] for i in range(0, len(metin), PAKET_KARAKTER)] or [""]
        for j, p in enumerate(parcalar, 1):
            if simdiki and boy + len(p) > PAKET_KARAKTER:
                paketler.append(simdiki)
                simdiki, boy = [], 0
            simdiki.append((Dosya(d.ad if len(parcalar) == 1 else f"{d.ad} (bölüm {j}/{len(parcalar)})", d.tur, ""), p))
            boy += len(p)
    if simdiki:
        paketler.append(simdiki)
    return paketler


def ozetle_model(dosyalar: list[Dosya], kod_bulgulari: list, terimler: list[str]) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    paketler = paketle(dosyalar, terimler)
    sonuc = {k: [] for k in LISTELER} | {"aciklama_talebi_sorulari": [], "genel_ozet": []}
    kod = "\n".join(f"- [{o}] {k}: {a}" for o, k, a in kod_bulgulari) or "- (yok)"
    for i, paket in enumerate(paketler, 1):
        ust = f"<paket no=\"{i}\" toplam=\"{len(paketler)}\"/>\n" if len(paketler) > 1 else ""
        mesaj = ust + "\n".join(f'<belge ad="{d.ad}" tur="{d.tur}">\n{m}\n</belge>' for d, m in paket) + f"\n<kod_kontrolleri>\n{kod}\n</kod_kontrolleri>"
        y = llm.json_iste(sistem, mesaj, SEMA)
        for k in LISTELER:
            sonuc[k] += [x for x in y.get(k, []) if isinstance(x, dict)]
        sonuc["aciklama_talebi_sorulari"] += [s for s in y.get("aciklama_talebi_sorulari", []) if s]
        if y.get("genel_ozet"):
            sonuc["genel_ozet"].append(y["genel_ozet"])
    gorulen, bilgiler = set(), []
    for b in sonuc["ihale_bilgileri"]:                         # paketler arası: aynı alanın ilk dolu değeri
        if b.get("alan") not in gorulen and str(b.get("deger", "")).strip():
            gorulen.add(b["alan"])
            bilgiler.append(b)
    sonuc["ihale_bilgileri"] = sorted(bilgiler, key=lambda b: BILGI_ALANLARI.index(b["alan"]) if b["alan"] in BILGI_ALANLARI else 99)
    sonuc["genel_ozet"] = "\n\n".join(sonuc["genel_ozet"])
    return sonuc


# ----------------------------------------------------------------------------
# Doğrulama ve hesap (kod)
# ----------------------------------------------------------------------------

def alinti_dogrula(kalem: dict, dosyalar: list[Dosya], maskeli: dict[str, str]) -> str:
    a = duz(kalem.get("alinti", ""))
    if len(a) < 8:
        return "Alıntı yok"
    kaynak = katla(kalem.get("kaynak", ""))
    adaylar = [d for d in dosyalar if katla(Path(d.ad).stem) in kaynak or katla(d.ad) in kaynak]
    for d in adaylar:
        if a in duz(maskeli[d.ad]):
            return "Doğrulandı"
    for d in dosyalar:
        if a in duz(maskeli[d.ad]):
            return f"Doğrulandı ({d.ad})" if adaylar else "Doğrulandı"
    return "Alıntı belgede bulunamadı"


def oran_dogrula(kalem: dict, alan: str) -> str:
    oran = kalem.get(alan)
    if oran is None:
        return ""
    try:
        o = Decimal(str(oran)).normalize()
    except InvalidOperation:
        return "Oran sayı değil"
    return "" if o in {x.normalize() for x in sayilar(kalem.get("alinti", ""))} else f"Oran ({o}) alıntıda rakamla geçmiyor"


def dogrula(sonuc: dict, dosyalar: list[Dosya], terimler: list[str], tarama: dict[str, list[str]]) -> list[tuple[str, str, str]]:
    maskeli = {d.ad: gizle(d.metin, terimler) for d in dosyalar}
    bulgular = []
    for liste in LISTELER:
        for k in sonuc[liste]:
            k["dogrulama"] = alinti_dogrula(k, dosyalar, maskeli)
            oran = oran_dogrula(k, "oran" if liste == "cezalar" else "oran_yuzde")
            if oran:
                k["dogrulama"] += f" · {oran}"
            if not k["dogrulama"].startswith("Doğrulandı") or oran:
                ad = k.get("alan") or k.get("kriter") or k.get("tur") or k.get("olay") or k.get("belge") or k.get("konu")
                bulgular.append(("Dikkat", ad, f"{k['dogrulama']} — belgede elle kontrol edin (kaynak: {k.get('kaynak') or '-'})"))
    ozet = katla(" ".join(str(v) for liste in LISTELER for k in sonuc[liste] for v in k.values()))
    for konu, yerler in tarama.items():
        anahtarlar = dict(KONULAR)[konu]
        if not any(a in ozet for a in anahtarlar):
            bulgular.append(("Dikkat", konu, "Belgede geçiyor ama özette yok: " + " | ".join(yerler[:2])))
    return bulgular


def hesapla(sonuc: dict, teklif: Decimal | None, yaklasik: Decimal | None) -> list[list]:
    """--teklif / --yaklasik verildiyse oranlardan tutar hesaplar: [kalem, oran, taban, tutar, açıklama]"""
    satirlar = []

    def taban_tutar(taban: str):
        if taban in ("teklif bedeli", "sözleşme bedeli"):
            return teklif, "teklif bedeli" if taban == "teklif bedeli" else "sözleşme bedeli (= teklif, ihale kazanılırsa)"
        if taban == "yaklaşık maliyet":
            return yaklasik, "yaklaşık maliyet"
        return None, taban

    kalemler = [(f"{t['tur']}", t.get("oran_yuzde"), Decimal(100), t.get("oran_tabani"), "en az" if t["tur"] == "Geçici teminat" else "")
                for t in sonuc["teminatlar"]]
    kalemler += [(f"{k['grup']}: {k['kriter']}", k.get("oran_yuzde"), Decimal(100), k.get("oran_tabani"), "alt sınır")
                 for k in sonuc["yeterlik_kriterleri"]]
    kalemler += [(f"{c['tur']} ({c['periyot']})", c.get("oran"), BIRIMLER.get(c.get("birim")), c.get("oran_tabani"), "")
                 for c in sonuc["cezalar"]]
    for ad, oran, bolen, taban, not_ in kalemler:
        if oran is None or bolen is None:
            continue
        tutar_tabani, taban_ad = taban_tutar(taban or "yok")
        o = Decimal(str(oran))
        birim = {100: "%", 1000: "‰", 10000: "‱"}[int(bolen)]
        if tutar_tabani is None:
            satirlar.append([ad, f"{o.normalize()}{birim}", taban_ad, None,
                             "Tutar için --teklif" + (" / --yaklasik" if taban == "yaklaşık maliyet" else "") + " verin"
                             if taban in TABANLAR[:3] else "Taban belirsiz"])
        else:
            satirlar.append([ad, f"{o.normalize()}{birim}", taban_ad, (tutar_tabani * o / bolen).quantize(Decimal("0.01")), not_])
    return satirlar


def takvim(sonuc: dict, dosyalar: list[Dosya], bugun: date) -> list[list]:
    ihale = ihale_tarihi_bul(dosyalar)
    if ihale is None:
        for b in sonuc["ihale_bilgileri"]:
            if b["alan"] == "İhale tarihi ve saati":
                m = re.search(r"(\d{2}\.\d{2}\.\d{4})", b["deger"])
                ihale = datetime.strptime(m.group(1), "%d.%m.%Y") if m else None
    if ihale is None:
        return [["İhale tarihi", "bulunamadı", "", "Belgeden elle girin"]]
    kalan = (ihale.date() - bugun).days
    satirlar = [["İhale tarihi", f"{ihale:%d.%m.%Y %H:%M}", kalan, "GEÇMİŞ" if kalan < 0 else ("Son hafta" if kalan <= 7 else "")]]
    for s in sonuc["sureler"]:
        m = re.search(r"ihale tarihinden\s*(\d+)\s*(?:\([^)]*\)\s*)?(?:takvim\s*)?gün\w*\s*önce", s.get("alinti", "") + " " + s.get("deger", ""), re.I)
        if m:
            t = ihale.date() - timedelta(days=int(m.group(1)))
            satirlar.append([f"{s['olay']} (son gün, hesap)", f"{t:%d.%m.%Y}", (t - bugun).days,
                             "GEÇMİŞ" if t < bugun else "Gün sayımı ve tatil kurallarını şartnameden doğrulayın"])
    return satirlar


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
ONAY_DOLGU = PatternFill("solid", fgColor="FFF4CE")
RENK = {"Hata": "FDE2E1", "Dikkat": "FFF4CE", "Bilgi": "E8F0FE", "yuksek": "FDE2E1", "orta": "FFF4CE", "dusuk": "E3F5E1"}
UST = Alignment(vertical="top", wrap_text=True)
KALIN = Font(bold=True)
RISK_AD = {"yuksek": "Yüksek", "orta": "Orta", "dusuk": "Düşük"}


def _sayfa(wb, ad, basliklar, genislik, satirlar, ai_sutunlar=(), onay=False, renk_sutun=None):
    ws = wb.create_sheet(ad)
    ws.append(basliklar + (["Kontrol / Not"] if onay else []))
    for h in ws[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for j, w in enumerate(list(genislik) + ([24] if onay else []), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for s in satirlar:
        ws.append(list(s) + ([""] if onay else []))
        for c in ai_sutunlar:
            ws.cell(ws.max_row, c).fill = AI_DOLGU
        if renk_sutun and str(s[renk_sutun - 1]) in RENK:
            ws.cell(ws.max_row, renk_sutun).fill = PatternFill("solid", fgColor=RENK[str(s[renk_sutun - 1])])
        if onay:
            ws.cell(ws.max_row, len(basliklar) + 1).fill = ONAY_DOLGU
        for h in ws[ws.max_row]:
            h.alignment = UST
            if isinstance(h.value, str) and h.value.startswith("Alıntı belgede bulunamadı"):
                h.fill = PatternFill("solid", fgColor=RENK["Hata"])
    ws.freeze_panes = "A2"
    return ws


def rapor_yaz(cikti: Path, dosyalar: list[Dosya], atlanan: list[str], sonuc: dict, bulgular: list, hesap: list, tkv: list) -> Path:
    wb = Workbook()
    o = wb.active
    o.title = "Özet"
    o.column_dimensions["A"].width, o.column_dimensions["B"].width, o.column_dimensions["C"].width = 30, 70, 40
    o.append(["İHALE DOKÜMANI ÖZETİ (TASLAK)"])
    o["A1"].font = Font(bold=True, size=13)
    o.append(["Belgeler", "; ".join(f"{d.ad} ({d.tur})" for d in dosyalar)])
    if atlanan:
        o.append(["Okunamayan", "; ".join(atlanan)])
    o.append([])
    o.append(["Alan", "Değer (model)", "Kaynak · Doğrulama"])
    for h in o[o.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for b in sonuc["ihale_bilgileri"]:
        o.append([b["alan"], b["deger"], f"{b['kaynak']} · {b['dogrulama']}"])
        o.cell(o.max_row, 2).fill = AI_DOLGU
    o.append([])
    o.append(["Takvim", "Tarih", "Kalan gün", "Not"])
    for h in o[o.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for t in tkv:
        o.append(t)
    o.append([])
    o.append(["Genel özet (model)", sonuc["genel_ozet"]])
    o.cell(o.max_row, 2).fill = AI_DOLGU
    o.append(["Model", llm.kullanim_ozeti()])
    o.append(["Not", "Taslaktır. Her kalemin kaynağı ve alıntısı verilmiştir; alıntılar kodla belgede aranmıştır. Teklif kararından önce "
                     "tüm doküman (zeyilnameler dahil) ve ilgili mevzuat okunmalıdır."])
    for r in o.iter_rows():
        for h in r:
            h.alignment = UST

    _sayfa(wb, "Yeterlik Kriterleri", ["Grup", "Kriter", "Açıklama", "Oran", "Kaynak", "Alıntı", "Doğrulama", "Firmamız Karşılıyor mu?"],
           (20, 30, 50, 8, 22, 50, 22, 18),
           [[k["grup"], k["kriter"], k["aciklama"], None if k.get("oran_yuzde") is None else f"%{Decimal(str(k['oran_yuzde'])).normalize()}",
             k["kaynak"], k["alinti"], k["dogrulama"], ""] for k in sonuc["yeterlik_kriterleri"]], ai_sutunlar=(1, 2, 3, 4))
    for r in wb["Yeterlik Kriterleri"].iter_rows(min_row=2):
        r[7].fill = ONAY_DOLGU
    _sayfa(wb, "Teminat ve Cezalar", ["Kalem", "Açıklama", "Kaynak", "Alıntı", "Doğrulama"], (26, 50, 22, 50, 24),
           [[t["tur"], t["aciklama"], t["kaynak"], t["alinti"], t["dogrulama"]] for t in sonuc["teminatlar"]] +
           [[c["tur"], c["aciklama"], c["kaynak"], c["alinti"], c["dogrulama"]] for c in sonuc["cezalar"]], ai_sutunlar=(1, 2))
    h = _sayfa(wb, "Hesaplar", ["Kalem", "Oran", "Taban", "Tutar (TL)", "Not"], (44, 10, 30, 18, 40), hesap)
    for r in h.iter_rows(min_row=2):
        r[3].number_format = "#,##0.00"
    _sayfa(wb, "Süreler", ["Olay", "Süre / tarih", "Kaynak", "Alıntı", "Doğrulama"], (34, 30, 22, 50, 24),
           [[s["olay"], s["deger"], s["kaynak"], s["alinti"], s["dogrulama"]] for s in sonuc["sureler"]], ai_sutunlar=(1, 2))
    _sayfa(wb, "Özel Şartlar ve Riskler", ["Risk", "Konu", "Açıklama", "Neden önemli", "Kaynak", "Alıntı", "Doğrulama"],
           (9, 26, 46, 40, 22, 46, 22),
           [[RISK_AD.get(x["risk"], x["risk"]), x["konu"], x["aciklama"], x["neden"], x["kaynak"], x["alinti"], x["dogrulama"]]
            for x in sorted(sonuc["ozel_sartlar"], key=lambda x: ["yuksek", "orta", "dusuk"].index(x["risk"]) if x["risk"] in RISK_AD else 3)],
           ai_sutunlar=(2, 3, 4), onay=True)
    for r in wb["Özel Şartlar ve Riskler"].iter_rows(min_row=2):
        ham = {v: k for k, v in RISK_AD.items()}.get(r[0].value)
        if ham:
            r[0].fill = PatternFill("solid", fgColor=RENK[ham])
    b = _sayfa(wb, "İstenen Belgeler", ["Belge", "Kaynak", "Doğrulama", "Hazır mı?", "Sorumlu"], (60, 22, 22, 12, 20),
               [[x["belge"], x["kaynak"], x["dogrulama"], "", ""] for x in sonuc["istenen_belgeler"]], ai_sutunlar=(1,))
    for r in b.iter_rows(min_row=2):
        r[3].fill = r[4].fill = ONAY_DOLGU
    _sayfa(wb, "Kontroller", ["Önem", "Konu", "Açıklama"], (10, 28, 110), bulgular, onay=True, renk_sutun=1)
    _sayfa(wb, "Açıklama Talepleri", ["No", "Soru taslağı (model)"], (6, 110),
           [[i, s] for i, s in enumerate(sonuc["aciklama_talebi_sorulari"], 1)], ai_sutunlar=(2,), onay=True)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

    md = cikti.with_suffix(".md")
    sat = ["# İhale Dokümanı Özeti (taslak)", "", "| Alan | Değer | Kaynak |", "|---|---|---|"]
    sat += [f"| {x['alan']} | {x['deger']} | {x['kaynak']} |" for x in sonuc["ihale_bilgileri"]]
    sat += ["", "## Takvim", ""] + [f"- {t[0]}: {t[1]} ({t[2]} gün) {t[3]}".rstrip() for t in tkv]
    sat += ["", "## Kod kontrolleri", ""] + ([f"- **{o_}** {k}: {a}" for o_, k, a in bulgular] or ["- Bulgu yok"])
    sat += ["", "## Yeterlik kriterleri", ""] + [f"- {k['grup']} — {k['kriter']}: {k['aciklama']} ({k['kaynak']})" for k in sonuc["yeterlik_kriterleri"]]
    if hesap:
        sat += ["", "## Hesaplar", ""] + [f"- {x[0]}: {x[1]} × {x[2]} = {tl(x[3]) if x[3] is not None else '-'} {x[4]}".rstrip() for x in hesap]
    sat += ["", "## Özel şartlar ve riskler", ""] + [f"- **{RISK_AD.get(x['risk'], x['risk'])}** {x['konu']}: {x['aciklama']} ({x['kaynak']})"
                                                   for x in sonuc["ozel_sartlar"]]
    sat += ["", "## Açıklama talebi soruları (taslak)", ""] + [f"{i}. {s}" for i, s in enumerate(sonuc["aciklama_talebi_sorulari"], 1)]
    sat += ["", "## Genel özet (model)", "", sonuc["genel_ozet"], "",
            "> Taslaktır. Teklif kararından önce tüm doküman, zeyilnameler ve mevzuat okunmalıdır."]
    md.write_text("\n".join(sat) + "\n", encoding="utf-8")
    return md


# ----------------------------------------------------------------------------
# Akış
# ----------------------------------------------------------------------------

def calistir(girdi: Path, cikti: Path, teklif: Decimal | None = None, yaklasik: Decimal | None = None, bugun: date | None = None,
             terimler: list[str] | None = None, evet: bool = False) -> dict:
    if not girdi.exists():
        raise llm.LLMHatasi(f"{girdi} bulunamadı.")
    dosyalar, atlanan = dosyalari_oku(girdi)
    if not dosyalar:
        raise llm.LLMHatasi(f"{girdi}: okunabilir belge yok. Atlanan: {', '.join(atlanan) or '-'}")
    terimler = terimler or []
    kod_bulgulari = tutarlilik(dosyalar)
    tarama = konu_taramasi(dosyalar)
    for d in dosyalar:
        print(f"[OK] {d.ad}: {d.tur} · {len(d.metin):,} karakter".replace(",", "."))
    for a in atlanan:
        print(f"[!] Atlandı: {a}")
    toplam = sum(len(d.metin) for d in dosyalar)
    llm.onay_al(f"{len(dosyalar)} belgenin metni ({toplam:,} karakter; telefon, e-posta ve {len(terimler)} gizli terim maskeli) "
                "gönderilecek.".replace(",", "."), evet)
    sonuc = ozetle_model(dosyalar, kod_bulgulari, terimler)
    bulgular = kod_bulgulari + dogrula(sonuc, dosyalar, terimler, tarama)
    hesap = hesapla(sonuc, teklif, yaklasik)
    tkv = takvim(sonuc, dosyalar, bugun or date.today())
    md = rapor_yaz(cikti, dosyalar, atlanan, sonuc, bulgular, hesap, tkv)
    return {"dosyalar": dosyalar, "sonuc": sonuc, "bulgular": bulgular, "hesap": hesap, "takvim": tkv, "md": md}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="İhale dokümanından yeterlik kriterleri, teminat, süre, ceza ve özel şartları özetler.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "pazar-yeri-ihalesi",
                   help="İhale dokümanı klasörü veya tek dosya (.pdf/.docx/.txt)")
    p.add_argument("--teklif", help="Planlanan teklif bedeli, TL (teminat, iş deneyimi ve ceza tutarlarını hesaplamak için)")
    p.add_argument("--yaklasik", help="Yaklaşık maliyet, TL (biliniyorsa)")
    p.add_argument("--bugun", help="Takvim hesabı için tarih GG.AA.YYYY (varsayılan: bugün)")
    p.add_argument("--gizle", nargs="*", default=[], metavar="TERİM", help="Modele gönderilmeden maskelenecek adlar")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "ihale_ozeti.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    teklif = tr_sayi(a.teklif) if a.teklif else None
    yaklasik = tr_sayi(a.yaklasik) if a.yaklasik else None
    if (a.teklif and teklif is None) or (a.yaklasik and yaklasik is None):
        print("[X] --teklif / --yaklasik sayı olmalı (ör. 48500000 veya 48.500.000)")
        return 2
    bugun = None
    if a.bugun:
        try:
            bugun = datetime.strptime(a.bugun, "%d.%m.%Y").date()
        except ValueError:
            print(f"[X] Tarih GG.AA.YYYY olmalı: {a.bugun}")
            return 2
    try:
        s = calistir(a.girdi, a.cikti, teklif, yaklasik, bugun, a.gizle, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    r = s["sonuc"]
    print(f"[OK] {len(r['yeterlik_kriterleri'])} yeterlik kriteri · {len(r['teminatlar'])} teminat · {len(r['cezalar'])} ceza · "
          f"{len(r['ozel_sartlar'])} özel şart · {len(r['istenen_belgeler'])} belge")
    for t in s["takvim"]:
        print(f"     {t[0]}: {t[1]} ({t[2]} gün) {t[3]}".rstrip())
    for o, k, a_ in s["bulgular"]:
        print(f"[{'X' if o == 'Hata' else '!'}] {k}: {a_}")
    print(f"[OK] Rapor: {a.cikti.resolve()} · {s['md'].resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
