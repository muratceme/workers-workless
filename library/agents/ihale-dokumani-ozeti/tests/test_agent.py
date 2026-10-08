"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

DOSYA = KLASOR / "ornek_veri" / "pazar-yeri-ihalesi"
I, S, T = "idari_sartname.txt", "sozlesme_tasarisi.txt", "teknik_sartname.txt"
MESAJLAR = []


def b(alan, deger, kaynak, alinti):
    return {"alan": alan, "deger": deger, "kaynak": kaynak, "alinti": alinti}


YANIT = {
    "ihale_bilgileri": [
        b("İhale kayıt numarası", "2026/999999", f"{I}, Madde 2.1", "İhale kayıt numarası: 2026/999999"),
        b("İdare", "Kurgusal Örnek Belediyesi Fen İşleri Müdürlüğü", f"{I}, Madde 1.1", "İdarenin adı: Kurgusal Örnek Belediyesi Fen İşleri Müdürlüğü"),
        b("İhale tarihi ve saati", "12.11.2026 10:30", f"{I}, Madde 3.2", "İhale tarihi ve saati: 12.11.2026 - 10:30"),
        b("Teklif türü", "Birim fiyat", f"{I}, Madde 4.1", "birim fiyat teklif şeklinde verecektir"),
        b("İşin süresi", "Yer tesliminden itibaren 300 takvim günü", f"{I}, Madde 6.1", "İşin süresi yer tesliminden itibaren 300 (üçyüz) takvim günüdür."),
        b("Fiyat farkı", "Verilmeyecek", f"{I}, Madde 14.1", "Bu ihalede fiyat farkı verilmeyecektir."),
        b("Avans", "Verilmeyecek", f"{I}, Madde 15.1", "Bu işte avans verilmeyecektir."),
        b("Alt yüklenici", "Öngörülmüyor", f"{I}, Madde 7.1", "alt yüklenici çalıştırılması öngörülmemektedir"),
        b("İhale kayıt numarası", "tekrar", f"{I}", "İhale kayıt numarası: 2026/999999"),
    ],
    "yeterlik_kriterleri": [
        {"grup": "İş deneyimi", "kriter": "İş deneyim belgesi", "aciklama": "Son 15 yılda benzer iş", "oran_yuzde": 50, "oran_tabani": "teklif bedeli",
         "kaynak": f"{I}, Madde 7.5.1", "alinti": "teklif edilen bedelin %50'sinden az olmamak üzere"},
        {"grup": "Ekonomik ve mali", "kriter": "Bilanço", "aciklama": "Cari oran ≥ 0,75; öz kaynak oranı ≥ 0,15; KV banka borcu/öz kaynak < 0,50",
         "oran_yuzde": None, "oran_tabani": "yok", "kaynak": f"{I}, Madde 7.4.1", "alinti": "cari oranın en az 0,75, öz kaynak oranının en az 0,15"},
        {"grup": "Ekonomik ve mali", "kriter": "İş hacmi — ciro", "aciklama": "Toplam ciro", "oran_yuzde": 25, "oran_tabani": "teklif bedeli",
         "kaynak": f"{I}, Madde 7.4.2", "alinti": "Toplam cironun teklif edilen bedelin %25'inden"},
        {"grup": "Ekonomik ve mali", "kriter": "İş hacmi — taahhüt", "aciklama": "Ciro yerine", "oran_yuzde": 15, "oran_tabani": "teklif bedeli",
         "kaynak": f"{I}, Madde 7.4.2", "alinti": "teklif edilen bedelin %15'inden az olmaması gerekir"},
        {"grup": "Teknik (personel, makine, kalite)", "kriter": "Anahtar teknik personel", "aciklama": "İnşaat, elektrik, makine mühendisi",
         "oran_yuzde": None, "oran_tabani": "yok", "kaynak": f"{I}, Madde 7.5.3", "alinti": "Anahtar teknik personel: 1 inşaat mühendisi"},
        {"grup": "Teknik (personel, makine, kalite)", "kriter": "ISO 9001", "aciklama": "Kalite yönetim sistemi belgesi", "oran_yuzde": None,
         "oran_tabani": "yok", "kaynak": f"{I}, Madde 7.5.2", "alinti": "Kalite yönetim sistemi belgesi: TS EN ISO 9001 belgesi."},
    ],
    "teminatlar": [
        {"tur": "Geçici teminat", "aciklama": "Teklif bedelinin en az %3'ü", "oran_yuzde": 3, "oran_tabani": "teklif bedeli",
         "kaynak": f"{I}, Madde 10.1", "alinti": "İstekliler teklif ettikleri bedelin %3'ünden az olmamak üzere"},
        {"tur": "Kesin teminat", "aciklama": "İhale bedelinin %6'sı", "oran_yuzde": 6, "oran_tabani": "sözleşme bedeli",
         "kaynak": f"{S}, Madde 4.1", "alinti": "sözleşme bedelinin %6'sı oranında kesin teminat"},
    ],
    "sureler": [
        {"olay": "Açıklama talebi", "deger": "İhale tarihinden 15 gün öncesine kadar", "kaynak": f"{I}, Madde 12.1",
         "alinti": "açıklama talepleri ihale tarihinden 15 gün öncesine kadar yazılı olarak yapılmalıdır"},
        {"olay": "Yer teslimi", "deger": "İmzadan itibaren 15 gün", "kaynak": f"{S}, Madde 3.1",
         "alinti": "imzalandığı tarihten itibaren 15 gün içinde yer teslimi yapılır"},        # uydurma: belgede 10 gün
        {"olay": "Teklif geçerlilik süresi", "deger": "90 takvim günü", "kaynak": f"{I}, Madde 11.1",
         "alinti": "Tekliflerin geçerlilik süresi, ihale tarihinden itibaren 90 (doksan) takvim günüdür."},
        {"olay": "Kesin kabul", "deger": "Geçici kabulden 12 ay sonra", "kaynak": f"{S}, Madde 9.1",
         "alinti": "Geçici kabul tarihinden itibaren 12 ay süre ile işin bakım ve onarımı yükleniciye aittir"},
    ],
    "cezalar": [
        {"tur": "Gecikme cezası", "aciklama": "Toplam %10'u aşarsa fesih", "oran": 0.6, "birim": "binde", "periyot": "günlük",
         "oran_tabani": "sözleşme bedeli", "kaynak": f"{S}, Madde 6.1", "alinti": "sözleşme bedelinin binde 0,6'sı oranında gecikme cezası uygulanır"},
    ],
    "istenen_belgeler": [
        {"belge": "İmza beyannamesi veya imza sirküleri", "kaynak": f"{I}, Madde 7.1.c",
         "alinti": "Teklif vermeye yetkili olduğunu gösteren imza beyannamesi veya imza sirküleri."},
    ],
    "ozel_sartlar": [
        {"konu": "Zemin riski", "aciklama": "Ek kazık ve iksa teklif fiyatına dahil", "risk": "yuksek", "neden": "Maliyeti belirsiz.",
         "kaynak": f"{T}, Madde 4.1", "alinti": "ek kazık ve iksa imalatları yüklenici tarafından teklif fiyatına dahil olarak yapılacaktır"},
        {"konu": "Pazar günleri", "aciklama": "Haftada iki gün çalışılamaz", "risk": "orta", "neden": "Süreyi uzatır.",
         "kaynak": f"{T}, Madde 2.1", "alinti": "pazar alanında çalışma yapmayacaktır"},
        {"konu": "İş artışı", "aciklama": "%10'a kadar", "risk": "dusuk", "neden": "Bilgi.", "kaynak": f"{S}, Madde 7.1",
         "alinti": "sözleşme bedelinin %10'una kadar iş artışı yapılabilir"},
        {"konu": "Hakediş", "aciklama": "Onaydan sonra 30 gün içinde ödeme", "risk": "orta", "neden": "Nakit akışı.", "kaynak": f"{S}, Madde 5.1",
         "alinti": "Hakediş raporu idarece onaylandıktan sonra 30 gün içinde ödeme yapılır."},
    ],
    "aciklama_talebi_sorulari": ["İşin süresi idari şartnamede 300, sözleşme tasarısında 330 takvim günüdür; hangisi esastır?"],
    "genel_ozet": "Kurgusal pazar yeri yapım işi.",
}


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    import copy
    return copy.deepcopy(YANIT)


class KodTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.d, cls.atlanan = agent.dosyalari_oku(DOSYA)

    def test_okuma(self):
        self.assertEqual([(x.ad, x.tur) for x in self.d], [(I, "İdari şartname"), (S, "Sözleşme tasarısı"), (T, "Teknik şartname")])
        self.assertEqual(self.atlanan, [])

    def test_tutarlilik(self):
        b = agent.tutarlilik(self.d)
        self.assertEqual(len(b), 1)                                        # gecikme cezası tek belgede, çelişki yok
        self.assertIn("300 takvim günü (idari_sartname.txt); 330 takvim günü (sozlesme_tasarisi.txt)", b[0][2])

    def test_ceza_celiskisi(self):
        ek = agent.Dosya("zeyilname.txt", "Zeyilname", "Gecikilen her takvim günü için sözleşme bedelinin on binde 5'i gecikme cezası kesilir.")
        b = agent.tutarlilik(self.d + [ek])
        self.assertIn("binde 0,6 (sozlesme_tasarisi.txt); on binde 5 (zeyilname.txt)", b[1][2])

    def test_ihale_tarihi_ve_tarama(self):
        self.assertEqual(f"{agent.ihale_tarihi_bul(self.d):%d.%m.%Y %H:%M}", "12.11.2026 10:30")
        t = agent.konu_taramasi(self.d)
        self.assertTrue(t["Sigorta"][0].startswith("sozlesme_tasarisi.txt: Yüklenici, iş yerini"))
        self.assertNotIn("Konsorsiyum", t)

    def test_sayilar(self):
        self.assertEqual(agent.sayilar("binde 0,6'sı ve 48.500.000 TL, %3"), {D("0.6"), D("48500000"), D("3")})
        self.assertEqual(agent.tr_sayi("48.500.000"), D("48500000"))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_uctan_uca(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "i.xlsx"
            s = agent.calistir(DOSYA, cikti, D("48500000"), None, date(2026, 10, 8), evet=True)
            m = MESAJLAR[0]
            self.assertNotIn("fenisleri@ornek.bel.tr", m)
            self.assertNotIn("0312 000 00 00", m)
            self.assertIn("<kod_kontrolleri>\n- [Hata] İşin süresi", m)
            r = s["sonuc"]
            self.assertEqual([x["alan"] for x in r["ihale_bilgileri"]][:3], ["İdare", "İhale kayıt numarası", "Teklif türü"])
            self.assertEqual(sum(1 for x in r["ihale_bilgileri"] if x["alan"] == "İhale kayıt numarası"), 1)
            self.assertEqual(r["ihale_bilgileri"][0]["dogrulama"], "Doğrulandı")
            yer = next(x for x in r["sureler"] if x["olay"] == "Yer teslimi")
            self.assertEqual(yer["dogrulama"], "Alıntı belgede bulunamadı")
            metin = "\n".join(f"{o} {k}: {a}" for o, k, a in s["bulgular"])
            self.assertIn("Hata İşin süresi", metin)
            self.assertIn("Dikkat Yer teslimi: Alıntı belgede bulunamadı", metin)
            self.assertIn("Dikkat Sigorta: Belgede geçiyor ama özette yok", metin)
            self.assertEqual(len(s["bulgular"]), 3)
            h = {x[0]: x[3] for x in s["hesap"]}
            self.assertEqual(h["Geçici teminat"], D("1455000.00"))
            self.assertEqual(h["Kesin teminat"], D("2910000.00"))
            self.assertEqual(h["İş deneyimi: İş deneyim belgesi"], D("24250000.00"))
            self.assertEqual(h["Ekonomik ve mali: İş hacmi — ciro"], D("12125000.00"))
            self.assertEqual(h["Ekonomik ve mali: İş hacmi — taahhüt"], D("7275000.00"))
            self.assertEqual(h["Gecikme cezası (günlük)"], D("29100.00"))
            self.assertEqual(s["takvim"], [["İhale tarihi", "12.11.2026 10:30", 35, ""],
                                           ["Açıklama talebi (son gün, hesap)", "28.10.2026", 20,
                                            "Gün sayımı ve tatil kurallarını şartnameden doğrulayın"]])
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Özet", "Yeterlik Kriterleri", "Teminat ve Cezalar", "Hesaplar", "Süreler", "Özel Şartlar ve Riskler",
                                             "İstenen Belgeler", "Kontroller", "Açıklama Talepleri"])
            self.assertEqual(wb["Özel Şartlar ve Riskler"]["A2"].value, "Yüksek")
            self.assertEqual(wb["Yeterlik Kriterleri"]["H1"].value, "Firmamız Karşılıyor mu?")
            md = s["md"].read_text(encoding="utf-8")
            self.assertIn("Gecikme cezası (günlük): 0.6‰ × sözleşme bedeli", md)
            self.assertIn("29.100,00 TL", md)

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_teklifsiz(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(DOSYA / I, Path(tmp) / "i.xlsx", bugun=date(2026, 11, 20), evet=True)
        self.assertTrue(all(x[3] is None and "--teklif" in x[4] for x in s["hesap"]))
        self.assertEqual(s["takvim"][0][3], "GEÇMİŞ")
        self.assertEqual([x.ad for x in s["dosyalar"]], [I])

    def test_oran_dogrula(self):
        self.assertEqual(agent.oran_dogrula({"oran_yuzde": 6, "alinti": "%6 oranında"}, "oran_yuzde"), "")
        self.assertIn("alıntıda rakamla geçmiyor", agent.oran_dogrula({"oran_yuzde": 5, "alinti": "%6 oranında"}, "oran_yuzde"))

    def test_klasor_yok(self):
        with self.assertRaises(llm.LLMHatasi):
            agent.calistir(Path("olmayan_klasor"), Path("x.xlsx"), evet=True)


if __name__ == "__main__":
    unittest.main()
