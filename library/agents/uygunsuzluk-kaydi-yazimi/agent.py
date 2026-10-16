"""
Uygunsuzluk Kaydı Yazımı — Workers / Workless AI Agent
Üretim › Kalite Kontrol Elemanı (tekstil, inşaat ve gıda kalite rolleri de kullanır)

Kalite kontrolcünün sahada aldığı kısa notları standart uygunsuzluk (NCR) kaydına çevirir.
1. Kod not dosyasını kayıtlara ayırır ("Kayıt: GG.AA.YYYY SS:DD | Yer" satırı), satırları numaralar; telefon, e-posta ve
   kimlik numaralarını maskeler. Hata kodu listesi (isteğe bağlı) okunur.
2. Model her kayıt için başlık, nesnel hata tanımı, ürün / malzeme, parti, kontrol / hatalı / etkilenen miktar, hata
   kodları, sınıf (Kritik / Majör / Minör), yapılan ilk müdahaleler, karar önerisi ve eksik bilgileri yazar; her kayıt
   dayandığı not satırlarını gösterir.
3. Kod kaydı denetler: sayılar, ürün ve parti notta geçiyor mu; hata kodu listede var mı; hatalı miktar kontrol miktarını
   aşıyor mu; notta gıda / insan güvenliği işareti varsa (metal, cam, yabancı madde, alerjen, yaralanma, taşıyıcı) sınıf
   Kritik mi. Sorun varsa yalnız o kayıt modele geri gönderilir (en fazla --tur kez).
4. Kod hata oranını, DÖF gerekip gerekmediğini (Kritik / Majör ya da aynı hata kodu birden çok kayıtta) ve sınıfa göre DÖF
   terminini hesaplar; eksik parti / miktar bilgisini işaretler.
Çıktının "DÖF Listesi" sayfası "DÖF Takibi" kod bloğunun girdi biçimindedir.

Kullanım:
    python agent.py                                              # örnek 4 kayıt
    python agent.py --girdi notlar.txt --kodlar hata_kodlari.csv --sorumlu "Kalite Mühendisi"
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import Counter
from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
RENK = {"Kritik": "FDE2E1", "Majör": "FFF4CE", "Minör": "E8F0FE"}
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}
VARSAYILAN_TERMIN = {"Kritik": 7, "Majör": 15, "Minör": 30}
# Bu ifadeler notta geçiyorsa sınıf Kritik olmalı (gıda güvenliği / insan güvenliği / yapısal güvenlik)
KRITIK = re.compile(r"\b(metal parca|metal dedektor|metal talas|metal bulas|cam kirig|cam parca|yabanci madde|alerjen|zehir|kontamin|yaralan|is kazasi|elektrik carp|yangin|patlama|"
                    r"tasiyici|yapisal|dokulme|cokme|catlak kolon|hayati|gida guvenligi|kimyasal bulas|kirik parca)")
BASLIK = re.compile(r"^\s*Kay[ıi]t\s*:\s*(\S+)(?:\s+(\d{1,2}:\d{2}))?\s*\|\s*(.+?)\s*$", re.I)


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1 or (s.count(".") == 1 and len(s.split(".")[1]) == 3):
        s = s.replace(".", "")
    try:
        return Decimal(s) if s not in ("", "-", ".") else None
    except InvalidOperation:
        return None


def sayilar(metin: str) -> set[Decimal]:
    return {v for x in re.findall(r"\d+(?:[.,]\d+)*", str(metin)) if (v := para(x)) is not None}


# ----------------------------------------------------------------------------
# Girdi
# ----------------------------------------------------------------------------

def kayitlari_oku(yol: Path) -> list[dict]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: not dosyası bulunamadı.")
    out = []
    for s in belge.metin_oku(yol).splitlines():
        m = BASLIK.match(s)
        if m:
            try:
                t = datetime.strptime(m[1] + (" " + m[2] if m[2] else ""), "%d.%m.%Y %H:%M" if m[2] else "%d.%m.%Y")
            except ValueError as h:
                raise llm.LLMHatasi(f"'{s.strip()}': tarih okunamadı (GG.AA.YYYY SS:DD).") from h
            out.append({"no": f"K{len(out) + 1}", "zaman": t, "yer": m[3], "satirlar": []})
        elif s.strip() and out:
            out[-1]["satirlar"].append(llm.maskele(s.strip()))
    out = [k for k in out if k["satirlar"]]
    if not out:
        raise llm.LLMHatasi(f"{yol.name}: 'Kayıt: GG.AA.YYYY SS:DD | Yer' başlıklı kayıt yok.")
    return out


def kodlari_oku(yol: Path | None) -> dict[str, str]:
    if yol is None:
        return {}
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: hata kodu listesi bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, read_only=True, data_only=True)
        satirlar = [[str(c or "").strip() for c in r] for r in wb.active.iter_rows(values_only=True)]
        wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=metin.split("\n")[0].count)))
    if not satirlar or len(satirlar[0]) < 2:
        raise llm.LLMHatasi(f"{yol.name}: 'Kod' ve 'Hata Tanımı' sütunları gerekli.")
    return {r[0].strip(): r[1].strip() for r in satirlar[1:] if len(r) >= 2 and r[0].strip()}


def termin_oku(metin: str) -> dict[str, int]:
    out = dict(VARSAYILAN_TERMIN)
    for parca in filter(None, (x.strip() for x in metin.split(","))):
        ad, _, gun = parca.partition("=")
        anahtar = {"kritik": "Kritik", "major": "Majör", "minor": "Minör"}.get(katla(ad))
        if not anahtar or not gun.strip().isdigit():
            raise ValueError(f"'{parca}' okunamadı (ör. kritik=7,majör=15,minör=30)")
        out[anahtar] = int(gun)
    return out


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

SINIF = ["Kritik", "Majör", "Minör"]
KARAR = ["Ayrıştır / beklet", "Yeniden işle", "Iskarta", "Şartlı kabul (yetkili onayıyla)", "Tedarikçiye iade", "Karar bekliyor"]
SEMA = {"type": "object", "properties": {
    "baslik": S, "hata_tanimi": S, "urun": S, "parti": S, "kontrol_miktari": S, "hatali_miktari": S, "etkilenen_miktar": S, "birim": S,
    "hata_kodlari": {"type": "array", "items": S}, "sinif": {"type": "string", "enum": SINIF},
    "ilk_mudahale": {"type": "array", "items": S}, "karar": {"type": "string", "enum": KARAR},
    "eksik_bilgiler": {"type": "array", "items": S}, "satirlar": {"type": "array", "items": {"type": "integer"}}},
    "required": ["baslik", "hata_tanimi", "urun", "parti", "kontrol_miktari", "hatali_miktari", "etkilenen_miktar", "birim", "hata_kodlari", "sinif",
                 "ilk_mudahale", "karar", "eksik_bilgiler", "satirlar"], "additionalProperties": False}
METIN_ALAN = {"baslik": "Başlık", "hata_tanimi": "Hata tanımı"}
SAYI_ALAN = {"kontrol_miktari": "Kontrol miktarı", "hatali_miktari": "Hatalı miktar", "etkilenen_miktar": "Etkilenen miktar"}


def denetle(k: dict, c: dict, kodlar: dict[str, str]) -> list[str]:
    on, notu = k["no"], "\n".join([k["yer"], *k["satirlar"]])      # kayıt başlığındaki yer de nottur
    sorun = []
    izinli = sayilar(notu) | {Decimal(k["zaman"].day), Decimal(k["zaman"].month), Decimal(k["zaman"].year)}
    if not c["satirlar"] or any(not 1 <= i <= len(k["satirlar"]) for i in c["satirlar"]):
        sorun.append(f"{on}: dayanak satırlar boş ya da geçersiz (1-{len(k['satirlar'])})")
    for a, ad in {**METIN_ALAN, **{"ilk": "İlk müdahale"}}.items():
        for m in (c["ilk_mudahale"] if a == "ilk" else [c[a]]):
            fazla = sorted(sayilar(m) - izinli, key=float)
            if fazla:
                sorun.append(f"{on}: {ad}: notta olmayan sayı: {', '.join(f'{x:g}' for x in fazla)}")
    for a, ad in SAYI_ALAN.items():
        v = c[a].strip()
        if v and (para(v) is None or para(v) not in izinli):
            sorun.append(f"{on}: {ad} '{v}' notta yazmıyor (hesaplamayın; notta yoksa boş bırakın)")
    for a, ad in (("urun", "Ürün"), ("parti", "Parti")):
        if c[a].strip() and not metinde_var(c[a], notu):
            sorun.append(f"{on}: {ad} '{c[a]}' notta aynen geçmiyor")
    h, kt = para(c["hatali_miktari"]), para(c["kontrol_miktari"])
    if h is not None and kt is not None and h > kt:
        sorun.append(f"{on}: hatalı miktar ({c['hatali_miktari']}) kontrol miktarından ({c['kontrol_miktari']}) büyük")
    if kodlar:
        yok = [x for x in c["hata_kodlari"] if x not in kodlar]
        if yok:
            sorun.append(f"{on}: hata kodu {', '.join(yok)} listede yok; listedeki kodları kullanın ya da boş bırakın")
    elif c["hata_kodlari"]:
        sorun.append(f"{on}: hata kodu listesi verilmedi; hata_kodlari boş olmalı")
    if (m := KRITIK.search(katla(notu + " " + k["yer"]))) and c["sinif"] != "Kritik":
        sorun.append(f"{on}: notta güvenlik işareti var ('{m.group(0)}'); sınıf Kritik olmalı")
    return sorun


def uret(kayitlar: list[dict], kodlar: dict[str, str], duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    liste = "\n".join(f"{a} | {b}" for a, b in kodlar.items()) or "(verilmedi: hata_kodlari boş kalsın)"
    sonuc = dict(onceki or {})
    hedef = kayitlar if duzeltme is None else [k for k in kayitlar if any(s.startswith(k["no"] + ":") for s in duzeltme)]
    for k in hedef:
        satir = "\n".join(f"{i}. {s}" for i, s in enumerate(k["satirlar"], 1))
        mesaj = f"<hata_kodlari>\n{liste}\n</hata_kodlari>\n<kayit tarih=\"{k['zaman']:%d.%m.%Y %H:%M}\" yer=\"{k['yer']}\">\n{satir}\n</kayit>"
        d = [s.split(": ", 1)[1] for s in (duzeltme or []) if s.startswith(k["no"] + ":")]
        if d:
            mesaj += ("\n<duzeltme>\nÖnceki kayıt: " + json.dumps(sonuc.get(k["no"]), ensure_ascii=False) + "\nKodun bulduğu sorunlar:\n"
                      + "\n".join(f"- {s}" for s in d) + "\n</duzeltme>")
        sonuc[k["no"]] = llm.json_iste(sistem, mesaj, SEMA)
    return sonuc


def duzeltme_dongusu(uret_f, denetle_f, tur: int) -> tuple[dict, list[str], int]:
    taslak = uret_f(None, None)
    kalan = denetle_f(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret_f(kalan, taslak)
        kalan = denetle_f(taslak)
    return taslak, kalan, yapilan


def degerlendir(kayitlar: list[dict], sonuc: dict, terminler: dict[str, int]) -> list[dict]:
    """Kod: hata oranı, eksik bilgi, DÖF gerekliliği ve termini."""
    sayac = Counter(x for k in kayitlar for x in set(sonuc[k["no"]]["hata_kodlari"]))
    out = []
    for k in kayitlar:
        c = sonuc[k["no"]]
        h, kt = para(c["hatali_miktari"]), para(c["kontrol_miktari"])
        oran = (h / kt * 100) if h is not None and kt else None
        eksik = list(c["eksik_bilgiler"])
        if not c["parti"].strip():
            eksik.insert(0, "Parti / lot / irsaliye no yok (izlenebilirlik)")
        if h is None and c["etkilenen_miktar"].strip() == "":
            eksik.insert(0, "Toplam hatalı / etkilenen miktar yazılmamış")
        tekrar = [x for x in c["hata_kodlari"] if sayac[x] > 1]
        dof = c["sinif"] in ("Kritik", "Majör") or bool(tekrar)
        neden = ([f"Sınıf {c['sinif']}"] if c["sinif"] != "Minör" else []) + ([f"Tekrarlayan hata kodu: {', '.join(tekrar)}"] if tekrar else [])
        out.append({"kayit": k, "c": c, "oran": oran, "eksik": list(dict.fromkeys(eksik)), "dof": dof, "dof_neden": "; ".join(neden),
                    "termin": (k["zaman"] + timedelta(days=terminler[c["sinif"]])).date() if dof else None})
    return out


# ----------------------------------------------------------------------------
# Çıktı
# ----------------------------------------------------------------------------

def _baslik(ws, basliklar, genislik) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font, h.alignment = BASLIK_DOLGU, BASLIK_YAZI, UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def rapor_yaz(cikti: Path, sonuclar: list[dict], kodlar: dict[str, str], kalan: list[str], tur: int, onek: str, sorumlu: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Uygunsuzluklar"
    _baslik(ws, ["Kayıt No", "Tarih", "Yer", "Başlık", "Hata Tanımı", "Ürün / Malzeme", "Parti", "Kontrol", "Hatalı", "Hata Oranı %", "Etkilenen", "Birim",
                 "Hata Kodları", "Sınıf", "İlk Müdahale", "Karar Önerisi", "Eksik Bilgiler", "DÖF", "DÖF Termini", "Dayanak Not", "Son Kontrol"],
            (11, 16, 26, 36, 50, 26, 14, 9, 9, 9, 10, 8, 26, 9, 45, 22, 36, 8, 11, 60, 30))
    for i, s in enumerate(sonuclar, 1):
        k, c = s["kayit"], s["c"]
        no = f"{onek}-{k['zaman']:%y%m%d}-{i:02d}"
        s["kayit_no"] = no
        sor = [x.split(": ", 1)[1] for x in kalan if x.startswith(k["no"] + ":")]
        ws.append([no, k["zaman"], k["yer"], c["baslik"], c["hata_tanimi"], c["urun"], c["parti"], _f(c["kontrol_miktari"]), _f(c["hatali_miktari"]),
                   None if s["oran"] is None else round(float(s["oran"]), 2), _f(c["etkilenen_miktar"]), c["birim"],
                   "\n".join(f"{x} {kodlar.get(x, '')}".strip() for x in c["hata_kodlari"]), c["sinif"], "\n".join(f"• {x}" for x in c["ilk_mudahale"]),
                   c["karar"], "\n".join(f"• {x}" for x in s["eksik"]), "Evet" if s["dof"] else "Hayır", s["termin"],
                   "\n".join(f"{j}. {k['satirlar'][j - 1]}" for j in c["satirlar"] if 1 <= j <= len(k["satirlar"])), "; ".join(sor) or "Geçti"])
        n = ws.max_row
        ws.cell(n, 2).number_format = "DD.MM.YYYY HH:MM"
        ws.cell(n, 19).number_format = "DD.MM.YYYY"
        for j in (4, 5, 6, 7, 8, 9, 11, 12, 13, 15, 16):
            ws.cell(n, j).fill = AI_DOLGU
        ws.cell(n, 14).fill = PatternFill("solid", fgColor=RENK[c["sinif"]])
        ws.cell(n, 21).fill = KIRMIZI if sor else YESIL
        for h in ws[n]:
            h.alignment = UST
    dof_sayfasi(wb.create_sheet("DÖF Listesi"), sonuclar, sorumlu)
    k = wb.create_sheet("Kontrol")
    k.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    k.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    k.cell(2, 1).fill = KIRMIZI if kalan else YESIL
    k.append(["Önemli: Mor hücreler yapay zekâ ile yazılmıştır. Sınıf ve karar önerisi, kaliteden sorumlu yetkilinin onayı olmadan uygulanmaz. "
              f"Hata oranı, DÖF gerekliliği ve termini kod tarafından hesaplanmıştır. DÖF listesi ayrıca {cikti.stem}_dof.xlsx dosyasındadır; "
              "o dosya 'DÖF Takibi' kod bloğuna (--doflar) doğrudan verilebilir."])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    ayri = Workbook()
    ayri.active.title = "DÖF Listesi"
    dof_sayfasi(ayri.active, sonuclar, sorumlu)
    ayri.save(cikti.with_name(cikti.stem + "_dof.xlsx"))


def dof_sayfasi(d, sonuclar: list[dict], sorumlu: str) -> None:
    """'DÖF Takibi' kod bloğunun girdi sütunları."""
    _baslik(d, ["DÖF No", "Tür", "Kaynak", "Konu", "Sorumlu", "Açılış Tarihi", "Termin", "Aksiyon Tamamlanma Tarihi", "Etkinlik Doğrulama Tarihi",
                "Etkinlik Sonucu"], (18, 10, 22, 50, 20, 12, 12, 14, 14, 12))
    for s in sonuclar:
        if s["dof"]:
            d.append([f"DOF-{s['kayit_no']}", "Düzeltici", f"Uygunsuzluk kaydı {s['kayit_no']}", s["c"]["baslik"], sorumlu, s["kayit"]["zaman"].date(),
                      s["termin"], None, None, None])
            d.cell(d.max_row, 6).number_format = d.cell(d.max_row, 7).number_format = "DD.MM.YYYY"


def _f(x):
    v = para(x)
    return None if v is None else float(v)


def calistir(girdi: Path, cikti: Path, kodlar_yolu: Path | None = None, sorumlu: str = "Kalite Sorumlusu", onek: str = "UGS",
             terminler: dict[str, int] | None = None, tur: int = 2, evet: bool = False) -> dict:
    kayitlar = kayitlari_oku(girdi)
    kodlar = kodlari_oku(kodlar_yolu)
    print(f"[OK] {len(kayitlar)} kayıt · {sum(len(k['satirlar']) for k in kayitlar)} not satırı · hata kodu: {len(kodlar)}")
    llm.onay_al(f"{len(kayitlar)} uygunsuzluk notu ve hata kodu listesi gönderilecek (telefon, e-posta, TCKN maskeli).", evet)
    sonuc, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(kayitlar, kodlar, d, o),
                                             lambda x: [s for k in kayitlar for s in denetle(k, x[k["no"]], kodlar)], tur)
    sonuclar = degerlendir(kayitlar, sonuc, terminler or VARSAYILAN_TERMIN)
    rapor_yaz(cikti, sonuclar, kodlar, kalan, yapilan, onek, sorumlu)
    return {"sonuclar": sonuclar, "kalan": kalan, "tur": yapilan}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Kalite kontrol saha notlarını standart uygunsuzluk kaydına ve DÖF listesine çevirir.")
    p.add_argument("--girdi", type=Path, default=ORNEK / "kalite_notlari.txt", help="Notlar (.txt/.docx); her kayıt 'Kayıt: GG.AA.YYYY SS:DD | Yer' ile başlar")
    p.add_argument("--kodlar", type=Path, default=None, help="Hata kodu listesi (.csv/.xlsx: Kod, Hata Tanımı)")
    p.add_argument("--sorumlu", default="Kalite Sorumlusu", help="DÖF listesindeki sorumlu")
    p.add_argument("--onek", default="UGS", help="Kayıt numarası öneki (UGS-261014-01)")
    p.add_argument("--termin", default="", help="Sınıfa göre DÖF termini, gün (varsayılan kritik=7,majör=15,minör=30)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "uygunsuzluklar.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        terminler = termin_oku(a.termin)
    except ValueError as h:
        print(f"[X] --termin {h}")
        return 2
    kodlar = a.kodlar or (ORNEK / "hata_kodlari.csv" if a.girdi == ORNEK / "kalite_notlari.txt" else None)
    try:
        s = calistir(a.girdi, a.cikti, kodlar, a.sorumlu, a.onek, terminler, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    say = Counter(x["c"]["sinif"] for x in s["sonuclar"])
    print(f"[OK] {len(s['sonuclar'])} kayıt · Kritik {say['Kritik']}, Majör {say['Majör']}, Minör {say['Minör']} · DÖF: {sum(x['dof'] for x in s['sonuclar'])}")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
