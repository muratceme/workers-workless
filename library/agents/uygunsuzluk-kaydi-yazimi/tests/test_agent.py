"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import json
import os
import sys
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
NOTLAR = ORNEK / "kalite_notlari.txt"
KODLAR = ORNEK / "hata_kodlari.csv"
CEVAP = json.loads((Path(__file__).parent / "sahte_cevaplar.json").read_text(encoding="utf-8"))
ANAHTAR = {"KP-50": "K1", "metal dedektör": "K2", "slump": "K3", "polo": "K4"}
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    """İlk turda K2 'Majör' (metal dedektör), K3 notta olmayan ürün adı, K4 hesaplanmış hatalı miktar (3 + 1 = 4) döner."""
    MESAJLAR.append(kullanici)
    kayit = kullanici.split("<kayit")[1]
    k = next(v for a, v in ANAHTAR.items() if a in kayit)
    tur = "duzeltme" if "<duzeltme>" in kullanici and k in CEVAP["duzeltme"] else "ilk"
    return copy.deepcopy(CEVAP[tur][k])


class KodTesti(unittest.TestCase):
    def test_okuma(self):
        k = agent.kayitlari_oku(NOTLAR)
        self.assertEqual([(x["no"], x["zaman"], x["yer"], len(x["satirlar"])) for x in k],
                         [("K1", datetime(2026, 10, 14, 9, 40), "Hat 2 - Kapak montaj", 4), ("K2", datetime(2026, 10, 14, 13, 15), "Dolum hattı - metal dedektör", 3),
                          ("K3", datetime(2026, 10, 15, 8, 20), "Şantiye B Blok 3. kat döşeme", 2), ("K4", datetime(2026, 10, 15, 10, 5), "Final kontrol", 2)])
        self.assertEqual(agent.kodlari_oku(KODLAR)["HK-07"], "Yabancı madde riski / metal dedektör hatası")
        self.assertEqual(agent.termin_oku("kritik=3, minör=45"), {"Kritik": 3, "Majör": 15, "Minör": 45})
        with self.assertRaises(ValueError):
            agent.termin_oku("acil=2")
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "n.txt"
            y.write_text("Kayıt: 32.10.2026 | Hat\nx\n", encoding="utf-8")
            with self.assertRaises(llm.LLMHatasi):
                agent.kayitlari_oku(y)
            y.write_text("Kayıt: 14.10.2026 | Hat 1\nÇapak var. Tel: 0532 111 22 33\n", encoding="utf-8")
            self.assertEqual(agent.kayitlari_oku(y)[0]["satirlar"], ["Çapak var. Tel: [TELEFON]"])

    def test_denetim(self):
        k = {x["no"]: x for x in agent.kayitlari_oku(NOTLAR)}
        kod = agent.kodlari_oku(KODLAR)
        self.assertEqual([agent.denetle(k[n], CEVAP["ilk"][n], kod) for n in ("K1",)], [[]])
        self.assertEqual([agent.denetle(k[n], CEVAP["duzeltme"][n], kod) for n in ("K2", "K3", "K4")], [[], [], []])
        self.assertEqual(agent.denetle(k["K2"], CEVAP["ilk"]["K2"], kod), ["K2: notta güvenlik işareti var ('metal dedektor'); sınıf Kritik olmalı"])
        self.assertEqual(agent.denetle(k["K3"], CEVAP["ilk"]["K3"], kod), ["K3: Ürün 'Hazır beton' notta aynen geçmiyor"])
        self.assertIn("Hatalı miktar '4' notta yazmıyor", agent.denetle(k["K4"], CEVAP["ilk"]["K4"], kod)[0])
        c = copy.deepcopy(CEVAP["ilk"]["K1"]) | {"hatali_miktari": "500", "kontrol_miktari": "12", "parti": "L2026-1013", "hata_kodlari": ["HK-99"],
                                                 "satirlar": [0, 9], "ilk_mudahale": ["Hat 7 saat durduruldu"]}
        s = agent.denetle(k["K1"], c, kod)
        self.assertEqual(len(s), 5)
        for parca in ("dayanak satırlar", "notta olmayan sayı: 7", "Parti 'L2026-1013'", "kontrol miktarından", "HK-99"):
            self.assertTrue(any(parca in x for x in s), parca)
        self.assertIn("hata kodu listesi verilmedi", agent.denetle(k["K1"], CEVAP["ilk"]["K1"], {})[0])
        self.assertFalse(agent.KRITIK.search(agent.katla("Metal kasada çizik var")))           # yalnız 'metal' Kritik saydırmaz
        self.assertTrue(agent.KRITIK.search(agent.katla("Ürün içinde cam kırığı bulundu")))

    def test_degerlendirme(self):
        k = agent.kayitlari_oku(NOTLAR)
        son = {**CEVAP["ilk"], **CEVAP["duzeltme"]}
        d = agent.degerlendir(k, son, agent.VARSAYILAN_TERMIN)
        self.assertEqual([(x["c"]["sinif"], x["dof"], x["termin"]) for x in d],
                         [("Majör", True, date(2026, 10, 29)), ("Kritik", True, date(2026, 10, 21)), ("Majör", True, date(2026, 10, 30)),
                          ("Majör", True, date(2026, 10, 30))])
        self.assertEqual(float(d[0]["oran"]), 2.4)
        self.assertEqual(d[1]["eksik"][0], "Toplam hatalı / etkilenen miktar yazılmamış")
        minor = copy.deepcopy(son)
        minor["K4"] |= {"sinif": "Minör"}
        minor["K1"] |= {"sinif": "Minör", "hata_kodlari": ["HK-21"]}
        d = agent.degerlendir(k, minor, agent.VARSAYILAN_TERMIN)
        self.assertEqual([(x["dof"], x["dof_neden"]) for x in (d[0], d[3])],
                         [(True, "Tekrarlayan hata kodu: HK-21"), (True, "Tekrarlayan hata kodu: HK-21")])
        minor["K1"]["hata_kodlari"] = ["HK-01"]
        self.assertEqual(agent.degerlendir(k, minor, agent.VARSAYILAN_TERMIN)[3]["dof"], False)
        k[0]["satirlar"][1] = "Ürün: 5 lt bidon kapağı KP-50."
        self.assertIn("Parti / lot / irsaliye no yok (izlenebilirlik)", agent.degerlendir(k, son | {"K1": son["K1"] | {"parti": ""}}, agent.VARSAYILAN_TERMIN)[0]["eksik"])


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ornek_ve_duzeltme(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "u.xlsx"
            self.assertEqual(agent.main(["--evet", "--sorumlu", "Kalite Mühendisi", "--cikti", str(cikti)]), 0)
            self.assertEqual(len(MESAJLAR), 7)                            # 4 kayıt + 3 düzeltme (K1 yeniden yazılmaz)
            self.assertIn("HK-07 | Yabancı madde riski", MESAJLAR[0])
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Uygunsuzluklar", "DÖF Listesi", "Kontrol"])
            ws = wb["Uygunsuzluklar"]
            self.assertEqual([ws.cell(r, 1).value for r in range(2, 6)], ["UGS-261014-01", "UGS-261014-02", "UGS-261015-03", "UGS-261015-04"])
            self.assertEqual([ws.cell(2, j).value for j in (7, 8, 9, 10, 14, 18)], ["L2026-1012", 500, 12, 2.4, "Majör", "Evet"])
            self.assertEqual(ws.cell(3, 14).value, "Kritik")
            self.assertEqual(ws.cell(5, 13).value, "HK-21 Leke / kirlilik\nHK-22 Dikiş hatası")
            self.assertTrue(all(ws.cell(r, 21).value == "Geçti" for r in range(2, 6)))
            d = load_workbook(Path(tmp) / "u_dof.xlsx").active
            self.assertEqual([c.value for c in d[1]][:7], ["DÖF No", "Tür", "Kaynak", "Konu", "Sorumlu", "Açılış Tarihi", "Termin"])
            self.assertEqual([d.cell(3, j).value for j in (1, 5)], ["DOF-UGS-261014-02", "Kalite Mühendisi"])
            self.assertEqual(d.cell(3, 7).value, datetime(2026, 10, 21))

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(NOTLAR, Path(tmp) / "u.xlsx", KODLAR, tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 3)
            self.assertIn("Kritik olmalı", load_workbook(Path(tmp) / "u.xlsx")["Uygunsuzluklar"].cell(3, 21).value)
            self.assertEqual(agent.main(["--termin", "acil=1", "--evet"]), 2)
            self.assertEqual(agent.main(["--girdi", str(Path(tmp) / "yok.txt"), "--evet"]), 1)
            self.assertEqual(agent.main(["--kodlar", str(Path(tmp) / "yok.csv"), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
