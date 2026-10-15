"""
Müşteri Şikâyeti Kök Neden Analizi — Workers / Workless AI Agent
Kalite Güvence ve Kalite Kontrol › Kalite Mühendisi

1. Kod şikâyetleri kümeler: aynı ürün + parti için --esik (varsayılan 2) ve üstü şikâyet "parti kümesi"; aynı üründe
   --urun-esik (varsayılan 3) ve üstü farklı partide şikâyet "ürün kümesi" (tasarım / süreç kaynaklı sistematik sorun
   şüphesi) olur.
2. Model her şikâyete hata türü verir ve her küme için 6M (insan, makine, malzeme, metot, ölçüm, çevre) kategorisinde
   olası kök neden hipotezleri, dayandığı şikâyetler, doğrulama adımı ve geçici önlem önerir.
3. Kod denetler: her şikâyet sınıflandırılmış mı ve alıntısı metinde mi, her kümenin en az bir hipotezi var mı,
   hipotezin dayandığı şikâyetler o kümeden mi, doğrulama adımı yazılmış mı, hipotez kesinmiş gibi ("kesinlikle",
   "kanıtlanmıştır") yazılmış mı. Sorun varsa modele geri gönderilir (en fazla --tur kez).
4. Rapor: Kümeler ve hipotezler (8D D4 için başlangıç), hata türü Pareto'su (şikâyet ve adet), ürün × parti matrisi.

Kullanım:
    python agent.py                                      # örnek şikâyetlerle dener
    python agent.py --girdi sikayetler.xlsx --esik 2 --urun-esik 3
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
HATA = ["Boyut / ölçü", "Fonksiyon (sızdırma, çalışmama)", "Görünüş (leke, çizik, renk)", "Malzeme / kırılma / çatlak", "Eksik adet / parça",
        "Etiket / baskı / kodlama", "Ambalaj / paketleme", "Kirlilik / yabancı madde", "Diğer"]
ALTI_M = ["İnsan", "Makine", "Malzeme", "Metot", "Ölçüm", "Çevre"]
KESIN = re.compile(r"\b(kesinlikle|kesin olarak|kanıtlanmıştır|kanıtlandı|şüphesiz|tartışmasız)\b", re.I)
SEMA = {
    "type": "object",
    "properties": {
        "sikayetler": {"type": "array", "items": {"type": "object", "properties": {"id": S, "hata_turu": {"type": "string", "enum": HATA}, "alinti": S},
                                                  "required": ["id", "hata_turu", "alinti"], "additionalProperties": False}},
        "kumeler": {"type": "array", "items": {"type": "object", "properties": {
            "kume": S, "ortak_bulgu": S,
            "hipotezler": {"type": "array", "items": {"type": "object", "properties": {
                "kategori": {"type": "string", "enum": ALTI_M}, "hipotez": S, "dayanak": {"type": "array", "items": S}, "dogrulama": S, "gecici_onlem": S},
                "required": ["kategori", "hipotez", "dayanak", "dogrulama", "gecici_onlem"], "additionalProperties": False}}},
            "required": ["kume", "ortak_bulgu", "hipotezler"], "additionalProperties": False}},
    },
    "required": ["sikayetler", "kumeler"],
    "additionalProperties": False,
}


def sikayetleri_oku(yol: Path) -> list[dict]:
    r = kayitlar(yol, {"id": ("Şikâyet No", "Şikayet No", "Kayıt No", "No"), "tarih": ("Tarih", "Şikâyet Tarihi", "Şikayet Tarihi"),
                       "musteri": ("Müşteri", "Müşteri Adı"), "urun": ("Ürün Kodu", "Ürün", "Malzeme Kodu"), "parti": ("Parti", "Lot", "Parti No", "Lot No"),
                       "hat": ("Hat", "Üretim Hattı"), "adet": ("Adet", "Etkilenen Adet", "Miktar"), "aciklama": ("Açıklama", "Şikâyet", "Şikayet", "Metin")},
                 ("aciklama", "urun"))
    out = []
    for i, x in enumerate(r, 1):
        if not str(x["aciklama"]).strip():
            continue
        try:
            adet = float(sayi_norm(x["adet"])) if str(x["adet"]).strip() else 0.0
        except ValueError:
            adet = 0.0
        out.append({"id": str(x["id"] or f"SK{i}"), "tarih": tarih_coz(x["tarih"]), "musteri": str(x["musteri"]), "urun": str(x["urun"]),
                    "parti": str(x["parti"]), "hat": str(x["hat"]), "adet": adet, "aciklama": str(x["aciklama"])})
    return out


def kumele(sik: list[dict], esik: int, urun_esik: int) -> list[dict]:
    parti, urun = defaultdict(list), defaultdict(list)
    for s in sik:
        if s["parti"]:
            parti[(s["urun"], s["parti"])].append(s["id"])
        urun[s["urun"]].append(s)
    out = []
    for (u, p), ids in sorted(parti.items(), key=lambda x: (-len(x[1]), x[0])):
        if len(ids) >= esik:
            out.append({"id": f"K{len(out) + 1}", "tur": "Parti", "urun": u, "parti": p, "sikayetler": ids})
    for u, l in sorted(urun.items()):
        partiler = sorted({s["parti"] for s in l if s["parti"]})
        if len(partiler) >= urun_esik:
            out.append({"id": f"K{len(out) + 1}", "tur": "Ürün (çok partili)", "urun": u, "parti": ", ".join(partiler), "sikayetler": [s["id"] for s in l]})
    return out


def denetle(t: dict, sik: list[dict], kumeler: list[dict]) -> list[str]:
    sorun = []
    sonuc = {x["id"]: x for x in t["sikayetler"]}
    for s in sik:
        x = sonuc.get(s["id"])
        if not x:
            sorun.append(f"{s['id']}: sınıflandırılmadı")
        elif not metinde_var(x["alinti"], s["aciklama"]):
            sorun.append(f"{s['id']}: alıntı açıklamada yok: '{x['alinti'][:50]}'")
    hip = {k["kume"]: k for k in t["kumeler"]}
    for k in kumeler:
        h = hip.get(k["id"])
        if not h or not h["hipotezler"]:
            sorun.append(f"{k['id']} ({k['urun']} / {k['parti']}): kök neden hipotezi yok")
            continue
        for x in h["hipotezler"]:
            disarida = [d for d in x["dayanak"] if d not in k["sikayetler"]]
            if not x["dayanak"]:
                sorun.append(f"{k['id']} · '{x['hipotez'][:40]}': dayanak şikâyet yok")
            if disarida:
                sorun.append(f"{k['id']} · '{x['hipotez'][:40]}': dayanak {disarida} bu kümede değil (kümedeki şikâyetler: {', '.join(k['sikayetler'])})")
            if not x["dogrulama"].strip():
                sorun.append(f"{k['id']} · '{x['hipotez'][:40]}': doğrulama adımı yok")
            if KESIN.search(x["hipotez"]):
                sorun.append(f"{k['id']} · '{x['hipotez'][:40]}': hipotez kesin sonuç gibi yazılmış; doğrulanana kadar 'olası' ifadesi kullanın")
    for k in hip:
        if k not in {x["id"] for x in kumeler}:
            sorun.append(f"Küme '{k}' kodun verdiği kümeler arasında yok")
    return sorun


def uret(sik: list[dict], kumeler: list[dict], duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    s = "\n".join(f'<sikayet id="{x["id"]}" urun="{x["urun"]}" parti="{x["parti"]}" hat="{x["hat"]}" tarih="{x["tarih"] or ""}" adet="{x["adet"]:g}">'
                  f"{llm.maskele(x['aciklama'])}</sikayet>" for x in sik)
    k = "\n".join(f'<kume id="{x["id"]}" tur="{x["tur"]}" urun="{x["urun"]}" parti="{x["parti"]}">{", ".join(x["sikayetler"])}</kume>' for x in kumeler)
    return llm.json_iste(sistem, f"<sikayetler>\n{s}\n</sikayetler>\n<kumeler>\n{k or 'yok'}\n</kumeler>" + duzeltme_blogu(duzeltme, onceki), SEMA)


def rapor_yaz(cikti: Path, sik, kumeler, t, kalan, tur) -> None:
    sonuc = {x["id"]: x for x in t["sikayetler"]}
    hip = {k["kume"]: k for k in t["kumeler"]}
    wb = Workbook()
    ws = wb.active
    ws.title = "Kümeler ve Hipotezler"
    baslik(ws, ["Küme", "Tür", "Ürün", "Parti", "Şikâyet Sayısı", "Etkilenen Adet", "Ortak Bulgu", "6M", "Olası Kök Neden", "Dayanak", "Doğrulama Adımı",
                "Geçici Önlem", "Sonuç (doğrulandı / elendi)"], (6, 14, 10, 22, 9, 10, 36, 9, 46, 18, 40, 34, 16))
    adet = {s["id"]: s["adet"] for s in sik}
    for k in kumeler:
        h = hip.get(k["id"], {"ortak_bulgu": "", "hipotezler": []})
        ilk = True
        for x in h["hipotezler"] or [None]:
            bas = [k["id"], k["tur"], k["urun"], k["parti"], len(k["sikayetler"]), sum(adet[i] for i in k["sikayetler"]), h["ortak_bulgu"]] if ilk else [""] * 7
            ws.append(bas + ([x["kategori"], x["hipotez"], ", ".join(x["dayanak"]), x["dogrulama"], x["gecici_onlem"], ""] if x else ["", "Hipotez yok"]))
            boya(ws, range(8, 13), AI_DOLGU)
            boya(ws, [13], SARI)
            ilk = False
    if not kumeler:
        ws.append(["—", "Eşiği aşan küme yok; şikâyetler tekil"])
    son_not(ws, "Önemli: Hipotezler yapay zekâ önerisidir ve doğrulanmamıştır. Kök neden, veri ve yerinde incelemeyle (8D D4) doğrulanmalıdır.", kalan, tur)
    hizala(ws)

    s = wb.create_sheet("Şikâyetler")
    kume_id = defaultdict(list)
    for k in kumeler:
        for i in k["sikayetler"]:
            kume_id[i].append(k["id"])
    baslik(s, ["Şikâyet No", "Tarih", "Müşteri", "Ürün", "Parti", "Hat", "Adet", "Açıklama", "Hata Türü", "Alıntı", "Küme"], (10, 11, 12, 10, 10, 8, 8, 60, 26, 30, 8))
    for x in sik:
        r = sonuc.get(x["id"], {})
        s.append([x["id"], x["tarih"], x["musteri"], x["urun"], x["parti"], x["hat"], x["adet"], x["aciklama"], r.get("hata_turu", "AI yanıt vermedi"),
                  r.get("alinti", ""), ", ".join(kume_id.get(x["id"], []))])
        s.cell(s.max_row, 2).number_format = "DD.MM.YYYY"
        boya(s, [9, 10], AI_DOLGU)
    s.freeze_panes = "B2"
    s.auto_filter.ref = f"A1:K{s.max_row}"
    hizala(s)

    p = wb.create_sheet("Pareto")
    baslik(p, ["Hata Türü", "Şikâyet", "Etkilenen Adet", "Kümülatif Şikâyet %"], (32, 9, 14, 18))
    say, ad = Counter(), Counter()
    for x in sik:
        if x["id"] in sonuc:
            say[sonuc[x["id"]]["hata_turu"]] += 1
            ad[sonuc[x["id"]]["hata_turu"]] += x["adet"]
    kum, top = 0, sum(say.values()) or 1
    for h, n in say.most_common():
        kum += n
        p.append([h, n, ad[h], round(100 * kum / top, 1)])
    if say:
        g = BarChart()
        g.title = "Hata türüne göre şikâyet"
        g.add_data(Reference(p, min_col=2, min_row=1, max_row=p.max_row), titles_from_data=True)
        g.set_categories(Reference(p, min_col=1, min_row=2, max_row=p.max_row))
        g.height, g.width = 8, 18
        p.add_chart(g, "F2")
    m = wb.create_sheet("Ürün × Parti")
    partiler = sorted({x["parti"] for x in sik})
    baslik(m, ["Ürün", *partiler, "Toplam"], (12, *[12] * len(partiler), 8))
    for u in sorted({x["urun"] for x in sik}):
        sat = [sum(1 for x in sik if x["urun"] == u and x["parti"] == pt) for pt in partiler]
        m.append([u, *[v or "" for v in sat], sum(sat)])
        for j, v in enumerate(sat, 2):
            if v >= 2:
                m.cell(m.max_row, j).fill = KIRMIZI
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(girdi: Path, cikti: Path, esik: int = 2, urun_esik: int = 3, tur: int = 2, evet: bool = False) -> dict:
    sik = sikayetleri_oku(girdi)
    if not sik:
        raise llm.LLMHatasi(f"{girdi.name}: şikâyet yok.")
    kumeler = kumele(sik, esik, urun_esik)
    print(f"[OK] {len(sik)} şikâyet · {len(kumeler)} küme")
    llm.onay_al(f"{len(sik)} şikâyetin ürün, parti, hat ve açıklaması gönderilecek (müşteri adı gönderilmez).", evet)
    t, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(sik, kumeler, d, o), lambda x: denetle(x, sik, kumeler), tur)
    rapor_yaz(cikti, sik, kumeler, t, kalan, yapilan)
    return {"sikayetler": sik, "kumeler": kumeler, "taslak": t, "kalan": kalan, "tur": yapilan}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    p = argparse.ArgumentParser(description="Müşteri şikâyetlerini ürün, parti ve hata türüne göre kümeleyip olası kök nedenleri önerir.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "sikayetler.csv",
                   help="Şikâyetler (.xlsx/.csv): Ürün Kodu, Açıklama; Şikâyet No, Tarih, Müşteri, Parti, Hat, Adet isteğe bağlı")
    p.add_argument("--esik", type=int, default=2, help="Aynı ürün + partide küme sayılacak en az şikâyet (varsayılan 2)")
    p.add_argument("--urun-esik", type=int, default=3, help="Ürün kümesi için şikâyet alınan en az farklı parti (varsayılan 3)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "sikayet_kok_neden.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.cikti, a.esik, a.urun_esik, a.tur, a.evet)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    hip = {k["kume"]: k for k in s["taslak"]["kumeler"]}
    for k in s["kumeler"]:
        print(f"[!] {k['id']} {k['tur']} {k['urun']} / {k['parti']}: {len(k['sikayetler'])} şikâyet · {len(hip.get(k['id'], {}).get('hipotezler', []))} hipotez")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
