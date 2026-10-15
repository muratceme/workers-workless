"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri" / "sikayetler.csv"
MESAJLAR = []
F, B, G, MA, E, ET = "Fonksiyon (sızdırma, çalışmama)", "Boyut / ölçü", "Görünüş (leke, çizik, renk)", "Malzeme / kırılma / çatlak", "Eksik adet / parça", "Etiket / baskı / kodlama"
SINIF = {"SK-501": (F, "vidalama sonrası sızıntı var"), "SK-502": (F, "Kapak dişleri eksik basılmış"), "SK-503": (B, "iç çapı ölçümde 37,6 mm çıktı"),
         "SK-504": (G, "siyah noktalar"), "SK-505": (ET, "mürekkep silik"), "SK-506": (ET, "Etiketler eğik yapıştırılmış"), "SK-507": (E, "2 adet eksik şişe"),
         "SK-508": (F, "kapak dönüyor ama sıkmıyor"), "SK-509": (F, "çapak ve eksik dolum"), "SK-510": (MA, "Şişe tabanında çatlak"),
         "SK-511": (F, "contası yerinden çıkmış")}


def h(kat, hip, dayanak, dog="Kalıp göz bazında ölçüm yapın"):
    return {"kategori": kat, "hipotez": hip, "dayanak": dayanak, "dogrulama": dog, "gecici_onlem": "Partiyi blokajda tutun"}


def taslak(dogru=True) -> dict:
    t = {"sikayetler": [{"id": i, "hata_turu": a, "alinti": b} for i, (a, b) in SINIF.items()],
         "kumeler": [
             {"kume": "K1", "ortak_bulgu": "Diş / çap sorunu ve sızdırma", "hipotezler": [
                 h("Makine", "Enjeksiyon basıncı / ütüleme yetersiz olabilir, dişler eksik dolmuş olabilir", ["SK-502", "SK-503"]),
                 h("Malzeme", "Conta tedarikçisinde ölçü sapması olabilir", ["SK-511"])]},
             {"kume": "K2", "ortak_bulgu": "Eksik dolum", "hipotezler": [h("Makine", "Kalıp gözlerinde tıkanma olabilir", ["SK-508", "SK-509"])]},
             {"kume": "K3", "ortak_bulgu": "Kodlama silik", "hipotezler": [h("Makine", "Mürekkep püskürtmeli kodlayıcı kafası kirli olabilir", ["SK-505", "SK-506"])]},
             {"kume": "K4", "ortak_bulgu": "Ambalaj", "hipotezler": [h("Metot", "Koli doldurma sayım kontrolü yetersiz olabilir", ["SK-507"])]},
             {"kume": "K5", "ortak_bulgu": "Üç partide kapak sorunu", "hipotezler": [h("Metot", "Kalıp bakım periyodu yetersiz olabilir", ["SK-501", "SK-508"])]}]}
    if not dogru:
        t["kumeler"][3]["hipotezler"] = []
        t["kumeler"][0]["hipotezler"][1]["dayanak"] = ["SK-511", "SK-504"]
        t["kumeler"][1]["hipotezler"][0]["hipotez"] = "Kesinlikle kalıp gözü tıkalı"
    return t


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return taslak("<duzeltme>" in kullanici)


class KodTesti(unittest.TestCase):
    def test_kumeleme(self):
        s = agent.sikayetleri_oku(ORNEK)
        k = agent.kumele(s, 2, 3)
        self.assertEqual([(x["id"], x["urun"], x["parti"], len(x["sikayetler"])) for x in k],
                         [("K1", "KPK-38", "L2609-14", 4), ("K2", "KPK-38", "L2610-05", 2), ("K3", "SSE-500", "L2609-30", 2),
                          ("K4", "SSE-500", "L2610-02", 2), ("K5", "KPK-38", "L2609-14, L2609-21, L2610-05", 7)])
        self.assertEqual(len(agent.kumele(s, 3, 4)), 1)

    def test_denetim(self):
        s = agent.sikayetleri_oku(ORNEK)
        k = agent.kumele(s, 2, 3)
        self.assertEqual(agent.denetle(taslak(), s, k), [])
        d = agent.denetle(taslak(False), s, k)
        self.assertEqual(len(d), 3)
        self.assertTrue(any("K4" in x and "hipotezi yok" in x for x in d) and any("['SK-504']" in x for x in d) and any("kesin sonuç" in x for x in d))
        t = taslak()
        t["sikayetler"].pop()
        t["sikayetler"][0]["alinti"] = "kapak kırık"
        t["kumeler"].append({"kume": "K9", "ortak_bulgu": "", "hipotezler": []})
        self.assertEqual(len(agent.denetle(t, s, k)), 3)


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "s.xlsx"
            s = agent.calistir(ORNEK, cikti, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            self.assertNotIn("Müşteri A", MESAJLAR[0])
            self.assertIn('<kume id="K5" tur="Ürün (çok partili)"', MESAJLAR[0])
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Kümeler ve Hipotezler", "Şikâyetler", "Pareto", "Ürün × Parti"])
            ws = wb["Kümeler ve Hipotezler"]
            self.assertEqual((ws.cell(2, 1).value, ws.cell(2, 6).value, ws.cell(3, 1).value), ("K1", 3350, None))
            self.assertEqual((wb["Pareto"].cell(2, 1).value, wb["Pareto"].cell(2, 2).value), (F, 5))

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(ORNEK, Path(tmp) / "s.xlsx", tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 3)
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--girdi", str(ORNEK.with_name("yok.csv")), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
