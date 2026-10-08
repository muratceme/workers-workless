"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    cevap = []
    for blok in re.findall(r'<soru id="([^"]+)".*?</soru>', kullanici, re.S):
        i = blok
        kaynaklar = re.findall(r'<kaynak id="(K\d+)"', kullanici.split(f'<soru id="{i}"')[1].split("</soru>")[0])
        c = {"id": i, "kapsam": "tam", "cevap": f"1. Adım [{kaynaklar[0]}]" if kaynaklar else "Bilgi yok.",
             "kaynaklar": kaynaklar[:1], "takip_sorusu": "", "insan_gerekli": False, "gerekce": ""}
        if i == "S-104":
            c["kaynaklar"] = ["K99"]                                  # uydurma kaynak → çıkarılmalı, kapsam yok
            c["cevap"] = "OneDrive kullanın."
        if i == "S-106":
            c["cevap"] = "Lütfen mevcut şifrenizi yazınız, kontrol edelim."  # parola isteyen taslak → reddedilmeli
        if i == "S-108":
            c.update(kapsam="yok", kaynaklar=[], cevap="Bilgi bankasında yok.")
        cevap.append(c)
    return {"cevaplar": cevap}


class AramaTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.b = agent.bilgi_bankasi_oku(ORNEK / "bilgi_bankasi")
        cls.s = agent.sorulari_oku(ORNEK / "sorular.csv")
        agent.on_isle(cls.s)
        agent.eslestir(cls.s, cls.b, 4, 1.0)
        cls.q = {x.id: x for x in cls.s}
        cls.ad = {x.id: x.baslik for x in cls.b}

    def test_bolumleme(self):
        self.assertEqual(len(self.b), 12)
        self.assertIn("Parola ve Hesap › Hesap kilitlendi", self.ad.values())

    def test_en_ilgili_bolum(self):
        beklenen = {"S-101": "VPN'e bağlanma", "S-102": "Parolamı unuttum", "S-103": "Yazıcı ekleme",
                    "S-104": "Posta kutusu boyutu", "S-109": "Telefona şirket e-postası"}
        for sid, baslik in beklenen.items():
            self.assertIn(baslik, self.ad[self.q[sid].eslesme[0][0]], sid)
        self.assertIn("Şüpheli e-posta", " ".join(self.ad[i] for i, _ in self.q["S-105"].eslesme[:2]))

    def test_etiketler_ve_parola(self):
        self.assertEqual(self.q["S-105"].etiketler, ["Güvenlik"])
        self.assertIn("Arıza olabilir", self.q["S-106"].etiketler)
        self.assertIn("Soruda parola var", self.q["S-102"].etiketler)
        self.assertNotIn("Soruda parola var", self.q["S-108"].etiketler)      # "Wi-Fi şifresi nedir?"
        self.assertEqual(agent.parola_maskele("şifrem Ankara2026 ama"), "şifrem [GİZLİ] ama")

    def test_basliksiz_metin(self):
        bolumler = agent.bolumle("a.txt", "Birinci paragraf.\n\nİkinci paragraf.")
        self.assertEqual(len(bolumler), 1)
        self.assertEqual(agent.kokler("Bağlanamıyorum VPN'e, lütfen"), ["bagla", "vpn"])


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_kod_denetimleri(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(agent.sorulari_oku(ORNEK / "sorular.csv"), ORNEK / "bilgi_bankasi", Path(tmp) / "c.xlsx", evet=True)
            wb = load_workbook(Path(tmp) / "c.xlsx")
            self.assertEqual(wb.sheetnames, ["Özet", "Cevaplar", "Kaynak Eşleşmeleri", "Bilgi Bankası Boşlukları"])
            bos = [r[0].value for r in wb["Bilgi Bankası Boşlukları"].iter_rows(min_row=3) if r[0].value and r[0].value.startswith("S-")]
            wb.close()
        r, n = s["sonuc"], s["notlar"]
        # gönderilen metin: parola maskeli, bilgi bankasının tamamı gönderilmez
        gonderilen = "".join(MESAJLAR)
        self.assertNotIn("Ankara2026", gonderilen)
        self.assertIn("[GİZLİ]", gonderilen)
        self.assertLess(gonderilen.count("<kaynak id="), 9 * 4 + 1)
        self.assertNotIn("sifre.ornek.local", MESAJLAR[0].split('<soru id="S-104"')[1].split("</soru>")[0])
        # uydurma kaynak
        self.assertEqual((r["S-104"]["kaynaklar"], r["S-104"]["kapsam"], r["S-104"]["insan_gerekli"]), ([], "yok", True))
        # parola isteyen taslak
        self.assertTrue(r["S-106"]["insan_gerekli"])
        self.assertTrue(any("parola istiyor" in x for x in n["S-106"]))
        # güvenlik kuralı
        self.assertTrue(r["S-105"]["insan_gerekli"])
        self.assertEqual(r["S-105"]["guven"], "orta")
        # güven ve boşluklar
        self.assertEqual(r["S-101"]["guven"], "yüksek")
        self.assertEqual(r["S-108"]["guven"], "düşük")
        self.assertEqual(set(bos), {"S-104", "S-108"})


if __name__ == "__main__":
    unittest.main()
