"""
İade EFT / FAST Müşteri Bilgilendirme — Workers / Workless AI Agent
Bankacılık Operasyon › Ödeme Sistemleri Operasyon Uzmanı

Alıcı bankadan iade gelen EFT / FAST ödemeleri için gönderen müşteriye bilgilendirme mesajı taslağı hazırlar.
1. Kod iade listesini okur ("İade Gelen EFT / FAST Analizi" kod bloğunun girdi biçimi) ve gönderen müşteri bazında toplar.
2. Modele yalnız iade kodu, banka açıklaması ve adet gider (müşteri no, alıcı adı, IBAN, tutar gönderilmez). Model her
   iade nedenini sade dille açıklar ve müşterinin yapması gerekeni yazar.
3. Kod açıklamayı denetler: her iade nedeni açıklanmış mı; metinde sayı, IBAN ya da yer tutucu var mı; ödeme / iade sözü
   ("hesabınıza yatırılmıştır", "yeniden gönderilecektir") var mı. Sorun varsa modele geri gönderilir (en fazla --tur kez).
4. Kod mesajı kurar: iade edilen ödemelerin listesi (tarih, tutar, sistem, alıcı, IBAN'ın son 4 hanesi, referans, neden),
   modelin açıklamaları, sizin verdiğiniz sabit bilgi cümlesi (--bilgi) ve imza.

Kullanım:
    python agent.py                                                    # örnek 4 iade, 3 müşteri
    python agent.py --iadeler iadeler.xlsx --bilgi "İade tutarı hesabınıza aktarılmıştır." --imza "Örnek Bank Ödeme Sistemleri"
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import llm

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}
SOZ = re.compile(r"\b(yatırılmıştır|yatırılacak\w*|aktarılmıştır|aktarılacak\w*|iade edilecek\w*|yeniden gönderilecek\w*|tekrar gönderilecek\w*|"
                 r"gönderilmiştir|garanti\w*|ücretsiz\w*|masrafsız\w*)\b", re.I)


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def para(x) -> Decimal | None:
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x or ""))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return Decimal(s) if s not in ("", "-", ".") else None
    except InvalidOperation:
        return None


def tl(x: Decimal) -> str:
    return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    m = re.search(r"\b(\d{1,2})\.(\d{1,2})\.(\d{4})\b", str(x or ""))
    try:
        return date(int(m[3]), int(m[2]), int(m[1])) if m else None
    except ValueError:
        return None


SUTUN = {"ref": ("referans", "islem no", "sorgu no", "referans no"), "gonderim": ("gonderim tarihi", "islem tarihi", "tarih"), "iade": ("iade tarihi",),
         "sistem": ("sistem", "odeme sistemi"), "musteri": ("gonderen musteri", "gonderen", "musteri", "musteri no"), "alici": ("alici", "alici adi"),
         "iban": ("alici iban", "iban"), "tutar": ("tutar",), "kod": ("iade kodu", "kod"), "aciklama": ("iade aciklamasi", "iade nedeni", "aciklama")}


def iadeleri_oku(yol: Path) -> list[dict]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: iade listesi bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, read_only=True, data_only=True)
        s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=metin.split("\n")[0].count)))
    s = [r for r in s if any(c not in (None, "") for c in r)]
    b = [katla(c) for c in s[0]] if s else []
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in SUTUN.items()}
    eksik = [SUTUN[a][0] for a in ("musteri", "tutar", "aciklama") if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun yok: {', '.join(eksik)}")
    out = []
    for r in s[1:]:
        al = lambda a: "" if idx[a] is None or idx[a] >= len(r) or r[idx[a]] is None else r[idx[a]]  # noqa: E731
        t = para(al("tutar"))
        if not str(al("musteri")).strip() or t is None:
            continue
        out.append({"ref": str(al("ref")).strip(), "gonderim": tarih_coz(al("gonderim")), "iade": tarih_coz(al("iade")), "sistem": str(al("sistem")).strip(),
                    "musteri": str(al("musteri")).strip(), "alici": str(al("alici")).strip(), "iban": re.sub(r"\s", "", str(al("iban"))), "tutar": t,
                    "kod": str(al("kod")).strip(), "aciklama": " ".join(str(al("aciklama")).split())})
    if not out:
        raise llm.LLMHatasi(f"{yol.name}: iade satırı yok.")
    return out


def nedenler(iadeler: list[dict]) -> list[dict]:
    """Aynı kod + açıklama (büyük / küçük harf farkı gözetilmeden) tek neden sayılır."""
    g: dict[tuple, dict] = {}
    for x in iadeler:
        k = (x["kod"], katla(x["aciklama"]))
        g.setdefault(k, {"anahtar": f"N{len(g) + 1}", "kod": x["kod"], "aciklama": x["aciklama"], "adet": 0})["adet"] += 1
    return list(g.values())


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

SEMA = {"type": "object", "properties": {"nedenler": {"type": "array", "items": {"type": "object", "properties": {
    "anahtar": S, "sade_aciklama": S, "musteri_ne_yapmali": S}, "required": ["anahtar", "sade_aciklama", "musteri_ne_yapmali"], "additionalProperties": False}}},
    "required": ["nedenler"], "additionalProperties": False}


def denetle(n: list[dict], c: dict) -> list[str]:
    sorun = []
    gelen = {x["anahtar"]: x for x in c["nedenler"]}
    for x in n:
        if x["anahtar"] not in gelen:
            sorun.append(f"{x['anahtar']} ({x['aciklama']}) açıklanmamış")
    for a, x in gelen.items():
        if a not in {y["anahtar"] for y in n}:
            sorun.append(f"{a}: listede olmayan neden")
            continue
        metin = x["sade_aciklama"] + " " + x["musteri_ne_yapmali"]
        if re.search(r"\d", metin):
            sorun.append(f"{a}: açıklamada sayı var; tutar, tarih, IBAN ve referansı kod ekler, siz yazmayın")
        if s := SOZ.search(metin):
            sorun.append(f"{a}: söz / kesinlik ifadesi ('{s.group(0)}'); bankanın yapacağı işlemi taahhüt etmeyin")
        if re.search(r"\[[A-ZÇĞİÖŞÜ-]+\]", metin):
            sorun.append(f"{a}: yer tutucu kullanılmış")
        if not x["musteri_ne_yapmali"].strip():
            sorun.append(f"{a}: müşterinin yapması gereken boş")
    return sorun


def uret(n: list[dict], duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    liste = "\n".join(f"{x['anahtar']} | kod: {x['kod'] or '-'} | banka açıklaması: {x['aciklama']} | {x['adet']} iade" for x in n)
    mesaj = f"<iade_nedenleri>\n{liste}\n</iade_nedenleri>"
    if duzeltme:
        mesaj += ("\n<duzeltme>\nÖnceki açıklamalar: " + json.dumps(onceki, ensure_ascii=False) + "\nKodun bulduğu sorunlar:\n"
                  + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")
    return llm.json_iste(sistem, mesaj, SEMA)


def duzeltme_dongusu(uret_f, denetle_f, tur: int) -> tuple[dict, list[str], int]:
    taslak = uret_f(None, None)
    kalan = denetle_f(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret_f(kalan, taslak)
        kalan = denetle_f(taslak)
    return taslak, kalan, yapilan


# ----------------------------------------------------------------------------
# Mesaj ve çıktı
# ----------------------------------------------------------------------------

def mesaj_kur(musteri_iadeleri: list[dict], n: list[dict], c: dict, bilgi: str, imza: str) -> str:
    acik = {x["anahtar"]: x for x in c["nedenler"]}
    anahtar = {(x["kod"], katla(x["aciklama"])): x["anahtar"] for x in n}
    l = ["Sayın Müşterimiz,", "", f"Aşağıdaki {'ödemeniz' if len(musteri_iadeleri) == 1 else 'ödemeleriniz'} alıcı bankadan iade edilmiştir:"]
    for x in musteri_iadeleri:
        iban = f"IBAN …{x['iban'][-4:]}" if len(x["iban"]) >= 4 else "IBAN -"
        l.append(f"- {x['gonderim']:%d.%m.%Y} tarihli {tl(x['tutar'])} TL {x['sistem']} · Alıcı: {x['alici'] or '-'} · {iban} · Ref: {x['ref'] or '-'} · "
                 f"İade nedeni: {x['aciklama']}" if x["gonderim"] else f"- {tl(x['tutar'])} TL · Alıcı: {x['alici']} · İade nedeni: {x['aciklama']}")
    l.append("")
    for k in dict.fromkeys(anahtar[(x["kod"], katla(x["aciklama"]))] for x in musteri_iadeleri):
        a = acik.get(k)
        if a:
            ad = next(y["aciklama"] for y in n if y["anahtar"] == k)
            l.append(f"{ad}: {a['sade_aciklama']} {a['musteri_ne_yapmali']}")
    if bilgi:
        l += ["", bilgi]
    l += ["", "Saygılarımızla,", imza]
    return "\n".join(l)


def rapor_yaz(cikti: Path, iadeler: list[dict], n: list[dict], c: dict, kalan: list[str], tur: int, bilgi: str, imza: str) -> Path:
    gruplar = defaultdict(list)
    for x in iadeler:
        gruplar[x["musteri"]].append(x)
    wb = Workbook()
    ws = wb.active
    ws.title = "Mesajlar"
    ws.append(["Gönderen Müşteri", "İade Adedi", "Toplam Tutar", "Mesaj"])
    for h in ws[1]:
        h.fill, h.font, h.alignment = BASLIK_DOLGU, BASLIK_YAZI, UST
    for j, w in enumerate((14, 8, 14, 120), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    metinler = []
    for m, xs in gruplar.items():
        t = mesaj_kur(xs, n, c, bilgi, imza)
        metinler.append(f"--- {m} ---\n{t}")
        ws.append([m, len(xs), float(sum((x["tutar"] for x in xs), Decimal(0))), t])
        ws.cell(ws.max_row, 3).number_format = "#,##0.00"
        ws.cell(ws.max_row, 4).alignment = UST
    a = wb.create_sheet("Neden Açıklamaları")
    a.append(["Kod", "Banka Açıklaması", "Adet", "Sade Açıklama", "Müşteri Ne Yapmalı"])
    for h in a[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    acik = {x["anahtar"]: x for x in c["nedenler"]}
    for x in n:
        y = acik.get(x["anahtar"], {"sade_aciklama": "", "musteri_ne_yapmali": ""})
        a.append([x["kod"], x["aciklama"], x["adet"], y["sade_aciklama"], y["musteri_ne_yapmali"]])
        a.cell(a.max_row, 4).fill = a.cell(a.max_row, 5).fill = AI_DOLGU
    for j, w in enumerate((7, 30, 6, 60, 60), 1):
        a.column_dimensions[get_column_letter(j)].width = w
    k = wb.create_sheet("Kontrol")
    k.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    k.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    k.cell(2, 1).fill = KIRMIZI if kalan else YESIL
    k.append(["Önemli: Ödeme listesi (tarih, tutar, alıcı, IBAN son 4 hane, referans) kod tarafından kurulmuştur; sağlayıcıya müşteri no, alıcı adı, "
              "IBAN ve tutar gönderilmemiştir. Sade açıklamalar yapay zekâ taslağıdır (mor); bankanızın iletişim standartlarına göre kontrol edin."])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    txt = cikti.with_suffix(".txt")
    txt.write_text("\n\n".join(metinler), encoding="utf-8")
    return txt


def calistir(girdi: Path, cikti: Path, bilgi: str = "", imza: str = "Ödeme Sistemleri Operasyon", tur: int = 2, evet: bool = False) -> dict:
    iadeler = iadeleri_oku(girdi)
    n = nedenler(iadeler)
    print(f"[OK] {len(iadeler)} iade · {len({x['musteri'] for x in iadeler})} müşteri · {len(n)} farklı neden")
    llm.onay_al(f"{len(n)} iade nedeninin kodu ve banka açıklaması gönderilecek. Müşteri no, alıcı adı, IBAN ve tutar gönderilmez.", evet)
    c, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(n, d, o), lambda x: denetle(n, x), tur)
    txt = rapor_yaz(cikti, iadeler, n, c, kalan, yapilan, bilgi, imza)
    return {"iadeler": iadeler, "nedenler": n, "c": c, "kalan": kalan, "tur": yapilan, "txt": txt}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="İade gelen EFT / FAST ödemeleri için gönderen müşteriye bilgilendirme mesajı taslağı hazırlar.")
    p.add_argument("--iadeler", type=Path, default=ORNEK / "iadeler.csv", help="İade listesi (.xlsx/.csv): Gönderen Müşteri, Alıcı, Alıcı IBAN, Tutar, İade Kodu, İade Açıklaması …")
    p.add_argument("--bilgi", default="", help="Mesaja eklenecek sabit bilgi cümlesi (ör. iade tutarının hesaba aktarımı hakkında bankanızın standart ifadesi)")
    p.add_argument("--imza", default="Ödeme Sistemleri Operasyon")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "iade_bilgilendirme.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.iadeler, a.cikti, a.bilgi, a.imza, a.tur, a.evet)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()} (+ {s['txt'].name})")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
