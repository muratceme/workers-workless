"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import json
import os
import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

IADE = KLASOR / "ornek_veri" / "iadeler.csv"
CEVAP = json.loads((Path(__file__).parent / "sahte_cevaplar.json").read_text(encoding="utf-8"))
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    """İlk turda limit aşımı için 'Ödemeniz yarın tekrar gönderilecektir' diye söz verir."""
    MESAJLAR.append(kullanici)
    return copy.deepcopy(CEVAP["dogru" if "<duzeltme>" in kullanici else "hatali"])


class KodTesti(unittest.TestCase):
    def test_okuma_ve_nedenler(self):
        i = agent.iadeleri_oku(IADE)
        self.assertEqual([(x["musteri"], x["tutar"], x["kod"]) for x in i][:2], [("M30316", D("1500.00"), "R02"), ("M30898", D("9800.00"), "R01")])
        n = agent.nedenler(i + [i[0] | {"aciklama": "HESAP KAPALI"}])            # büyük / küçük harf farkı aynı neden
        self.assertEqual([(x["anahtar"], x["kod"], x["adet"]) for x in n], [("N1", "R02", 2), ("N2", "R01", 1), ("N3", "R01", 1), ("N4", "R07", 1)])

    def test_denetim(self):
        n = agent.nedenler(agent.iadeleri_oku(IADE))
        self.assertEqual(agent.denetle(n, CEVAP["dogru"]), [])
        self.assertEqual(agent.denetle(n, CEVAP["hatali"]),
                         ["N4: söz / kesinlik ifadesi ('tekrar gönderilecektir'); bankanın yapacağı işlemi taahhüt etmeyin"])
        c = copy.deepcopy(CEVAP["dogru"])
        c["nedenler"] = c["nedenler"][:2] + [{"anahtar": "N9", "sade_aciklama": "x", "musteri_ne_yapmali": "y"},
                                             {"anahtar": "N3", "sade_aciklama": "1.500 TL [IBAN] iade", "musteri_ne_yapmali": ""}]
        s = agent.denetle(n, c)
        for parca in ("N4 (Limit aşımı) açıklanmamış", "N9: listede olmayan", "N3: açıklamada sayı", "N3: yer tutucu", "N3: müşterinin yapması gereken boş"):
            self.assertTrue(any(parca in x for x in s), parca)


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ornek_ve_mesaj(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "m.xlsx"
            self.assertEqual(agent.main(["--evet", "--bilgi", "Standart bilgi.", "--imza", "Örnek Bank", "--cikti", str(cikti)]), 0)
            self.assertEqual(len(MESAJLAR), 2)
            for m in MESAJLAR:                                               # müşteri, alıcı, IBAN, tutar gönderilmez
                for x in ("M30316", "Deneme Yapı", "TR07", "1.500", "REF700109"):
                    self.assertNotIn(x, m)
            ws = load_workbook(cikti)["Mesajlar"]
            self.assertEqual([ws.cell(r, 1).value for r in range(2, 5)], ["M30316", "M30898", "M30265"])
            t = ws.cell(2, 4).value
            self.assertIn("- 01.09.2026 tarihli 1.500,00 TL FAST · Alıcı: Deneme Yapı Ltd. · IBAN …2411 · Ref: REF700109 · İade nedeni: Hesap kapalı", t)
            self.assertTrue(t.endswith("Standart bilgi.\n\nSaygılarımızla,\nÖrnek Bank"))
            self.assertIn("ödemeleriniz", t)
            self.assertIn("Aşağıdaki ödemeniz", ws.cell(3, 4).value)

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(agent.main(["--iadeler", str(Path(tmp) / "yok.csv"), "--evet"]), 1)
            y = Path(tmp) / "i.csv"
            y.write_text("Referans;Tutar\nR1;100\n", encoding="utf-8")
            self.assertEqual(agent.main(["--iadeler", str(y), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
