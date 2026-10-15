"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

LISTE = KLASOR / "ornek_veri" / "dosyalar.csv"
BUGUN = date(2026, 10, 15)
MESAJLAR = []


def rapor(m, dogru=True) -> dict:
    if m == "M1":
        t = {"genel_ozet": "Sayın {MUVEKKIL}, büromuzda takip edilen 3 dosyanızın güncel durumu aşağıdadır. Dosyalarınızdaki toplam kalan alacak "
                           "1.745.000,00 TL'dir.",
             "dosyalar": [{"dosya_no": "2025/101", "durum": "{KT1} aleyhine açılan davada bilirkişi raporu (mahkemenin görüşüne başvurduğu uzmanın raporu) "
                                                           "tebliğ edildi ve rapora itiraz dilekçemiz verildi. Bir sonraki duruşma 19.11.2026 tarihindedir."},
                          {"dosya_no": "2025/102", "durum": "{KT2} hakkındaki icra takibinde 15.09.2026 tarihinde borçlunun banka hesaplarına haciz bildirimi "
                                                           "gönderildi. 380.000,00 TL alacağın 95.000,00 TL'si tahsil edildi; kalan alacak 285.000,00 TL'dir."},
                          {"dosya_no": "2024/087", "durum": "{KT3} hakkındaki takipte taşınmaz için kıymet takdiri (değer tespiti) 03.05.2026 tarihinde talep "
                                                           "edildi; o tarihten bu yana dosyada yeni işlem yok. Kalan alacak 210.000,00 TL'dir."}],
             "muvekkilden_beklenenler": ["{KT2} adına bildiğiniz başka banka hesabı ya da taşınmaz varsa bize iletmeniz"]}
        if not dogru:
            t["dosyalar"][0]["durum"] = "{KT1} aleyhine açılan davada bir sonraki duruşma 26.11.2026 tarihindedir. Davayı kazanacağımız kesindir."
            del t["dosyalar"][2]
        return t
    return {"genel_ozet": "Sayın {MUVEKKIL}, büromuzda takip edilen 1 dosyanızın durumu aşağıdadır.",
            "dosyalar": [{"dosya_no": "2026/015", "durum": "Ankara 2. İş Mahkemesindeki davada 25.09.2026 tarihinde ön inceleme duruşması yapıldı ve tanık "
                                                          "listesi istendi. Tanık listesinin 30.10.2026 tarihine kadar sunulması gerekiyor."}],
            "muvekkilden_beklenenler": ["Tanıklarınızın ad, adres ve iletişim bilgilerini büromuza iletmeniz"]}


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return rapor("M1" if "2025/101" in kullanici else "M2", "<duzeltme>" in kullanici)


def gruplar():
    return agent.muvekkiller(agent.dosyalari_oku(LISTE, BUGUN))


class KodTesti(unittest.TestCase):
    def test_hesap(self):
        g = gruplar()
        self.assertEqual([(x["kimlik"], x["ad"], len(x["dosyalar"])) for x in g], [("M1", "Örnek İnşaat A.Ş.", 3), ("M2", "Deneme Gıda Ltd. Şti.", 1)])
        d = g[0]["dosyalar"]
        self.assertEqual((d[1]["kalan_v"], d[2]["hareketsiz_gun"], d[0]["kalan_gun"]), (285000.0, 165, 35))
        self.assertEqual(g[0]["karsi"], {"Deneme Yapı Ltd. Şti.": "{KT1}", "Kaya Tesisat": "{KT2}", "Yıldız Nakliyat": "{KT3}"})
        self.assertIn("165 gündür işlem yok", agent.dosya_satiri(d[2], g[0]["karsi"], 90))

    def test_denetim(self):
        g = gruplar()
        self.assertEqual(agent.denetle(g, {"M1": rapor("M1"), "M2": rapor("M2")}, 90), [])
        s = agent.denetle(g, {"M1": rapor("M1", False), "M2": rapor("M2")}, 90)
        self.assertEqual(len(s), 4, s)
        self.assertTrue(all(x.startswith("M1:") for x in s) and any("2024/087" in x for x in s) and any("26.11.2026" in x for x in s)
                        and any("19.11.2026" in x for x in s) and any("kazanaca" in x for x in s))
        sizinti = rapor("M2")
        sizinti["dosyalar"][0]["durum"] += " Benzer bir dosyada (2025/8841) haciz uygulandı. {KT2} ile görüşüldü."
        s = agent.denetle(g, {"M1": rapor("M1"), "M2": sizinti}, 90)
        self.assertTrue(any("başka müvekkilin" in x and "2025/8841 E." in x for x in s) and any("{KT2}" in x for x in s))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "r.xlsx"
            s = agent.calistir(LISTE, cikti, "Örnek Hukuk Bürosu", BUGUN, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (3, [], 1))
            for gizli in ("Örnek İnşaat", "Deneme Yapı", "Kaya Tesisat", "Deneme Gıda"):
                self.assertNotIn(gizli, "".join(MESAJLAR))
            self.assertIn("2025/101", MESAJLAR[2])
            m = s["dosyalar"][0].read_text(encoding="utf-8")
            self.assertIn("Sayın Örnek İnşaat A.Ş., büromuzda", m)
            self.assertIn("Kaya Tesisat hakkındaki icra takibinde", m)
            self.assertIn("taahhüt içermez", m)
            ws = load_workbook(cikti)["Dosyalar"]
            self.assertEqual((ws.max_row, ws["G4"].value), (9, 165))

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(LISTE, Path(tmp) / "r.xlsx", "Büro", BUGUN, tur=0, evet=True)
            self.assertEqual((len(s["kalan"]), len(s["dosyalar"])), (4, 1))
            self.assertEqual(agent.main(["--evet", "--bugun", "15.10.2026", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--evet", "--bugun", "yarın"]), 1)


if __name__ == "__main__":
    unittest.main()
