"""
Müvekkil Dosya Durum Raporu — Workers / Workless AI Agent
Dava Takip › Avukat (Büro)

1. Kod dosya listesini okur (Dosya No; Müvekkil; Tür; Mahkeme / Daire; Esas No; Karşı Taraf; Son İşlem Tarihi; Son İşlem;
   Sonraki İşlem Tarihi; Sonraki İşlem; Alacak Tutarı; Tahsil Edilen), dosyaları müvekkile göre gruplar ve hesaplar:
   kalan alacak, son işlemden bu yana geçen gün (hareketsiz dosya), sonraki işleme kalan gün. Müvekkil ve karşı taraf adları
   modele gönderilmez; {MUVEKKIL}, {KT1}, {KT2} … kodlarıyla gönderilir.
2. Model her müvekkil için genel özet, dosya bazında sade dille durum açıklaması ve müvekkilden beklenenleri yazar.
3. Kod denetler:
   - Müvekkilin her dosyası raporda var mı; başka müvekkilin dosya / esas numarası rapora karışmış mı
   - Dosya metnindeki tarih ve tutarlar o dosyanın kayıtlarında (ya da hesaplanan kalan alacakta) var mı
   - Sonraki işlem tarihi dosya metninde aynen yazılmış mı
   - "Kazanacağız", "kesin", "garanti" gibi sonuç vaadi var mı; başka müvekkilin karşı taraf kodu kullanılmış mı
   Sorun varsa yalnız o müvekkilin raporu modele geri gönderilir (en fazla --tur kez).
4. Kod adları geri yazar ve müvekkil başına rapor dosyası üretir; raporlar gönderilmez.

Kullanım:
    python agent.py --bugun 15.10.2026                                 # örnek dosya listesi
    python agent.py --girdi dosyalar.xlsx --buro "Av. Ad Soyad Hukuk Bürosu" --hareketsiz 60
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
SEMA = {
    "type": "object",
    "properties": {"genel_ozet": S,
                   "dosyalar": {"type": "array", "items": {"type": "object", "properties": {"dosya_no": S, "durum": S},
                                                           "required": ["dosya_no", "durum"], "additionalProperties": False}},
                   "muvekkilden_beklenenler": {"type": "array", "items": S}},
    "required": ["genel_ozet", "dosyalar", "muvekkilden_beklenenler"],
    "additionalProperties": False,
}
ALANLAR = {"dosya": ("Dosya No", "Büro Dosya No"), "muvekkil": ("Müvekkil",), "tur": ("Tür", "Dosya Türü"), "mahkeme": ("Mahkeme / Daire", "Mahkeme", "İcra Dairesi"),
           "esas": ("Esas No", "Esas"), "karsi": ("Karşı Taraf",), "son_tarih": ("Son İşlem Tarihi",), "son": ("Son İşlem",),
           "sonraki_tarih": ("Sonraki İşlem Tarihi", "Duruşma Tarihi"), "sonraki": ("Sonraki İşlem",), "alacak": ("Alacak Tutarı", "Alacak"),
           "tahsil": ("Tahsil Edilen", "Tahsilat")}
VAAT = re.compile(r"\b(kazanaca\w*|kazanırız|kesinlikle|kesindir|kesin olarak|garanti\w*|mutlaka lehimize|kaybetme(?:miz|niz) mümkün değil|şüphesiz)\b", re.I)


def para(v) -> float | None:
    s = str(v if v is not None else "").strip()
    if not s:
        return None
    try:
        return float(v) if isinstance(v, (int, float)) else float(sayi_norm(s.replace("TL", "").strip()))
    except ValueError:
        return None


def tl(v: float | None) -> str:
    return "" if v is None else f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".") + " TL"


def dosyalari_oku(yol: Path, bugun: date) -> list[dict]:
    r = kayitlar(yol, ALANLAR, ("dosya", "muvekkil"))
    out = []
    for x in r:
        if not str(x["dosya"]).strip():
            continue
        d = {a: (str(v).strip() if not isinstance(v, (date, datetime)) else v) for a, v in x.items()}
        for a in ("son_tarih", "sonraki_tarih"):
            t = tarih_coz(x[a])
            d[a] = f"{t:%d.%m.%Y}" if t else str(x[a] or "").strip()
            d[a + "_d"] = t
        d["alacak_v"], d["tahsil_v"] = para(x["alacak"]), para(x["tahsil"])
        d["kalan_v"] = None if d["alacak_v"] is None else d["alacak_v"] - (d["tahsil_v"] or 0)
        d["hareketsiz_gun"] = (bugun - d["son_tarih_d"]).days if d["son_tarih_d"] else None
        d["kalan_gun"] = (d["sonraki_tarih_d"] - bugun).days if d["sonraki_tarih_d"] else None
        out.append(d)
    return out


def muvekkiller(dosyalar: list[dict]) -> list[dict]:
    m = {}
    for d in dosyalar:
        k = katla(d["muvekkil"])
        if k not in m:
            m[k] = {"kimlik": f"M{len(m) + 1}", "ad": d["muvekkil"], "dosyalar": [], "karsi": {}}
        g = m[k]
        g["dosyalar"].append(d)
        if d["karsi"] and d["karsi"] not in g["karsi"]:
            g["karsi"][d["karsi"]] = f"{{KT{len(g['karsi']) + 1}}}"
    return list(m.values())


def dosya_rakamlari(d: dict) -> set[str]:
    alanlar = [d[a] for a in ("dosya", "mahkeme", "esas", "son_tarih", "son", "sonraki_tarih", "sonraki", "alacak", "tahsil")]
    return set().union(*(sayilar(x) for x in alanlar), *(sayilar(tl(d[a])) for a in ("alacak_v", "tahsil_v", "kalan_v") if d[a] is not None))


def dosya_satiri(d: dict, karsi: dict, hareketsiz: int) -> str:
    h = d["hareketsiz_gun"]
    return (f"<dosya no=\"{d['dosya']}\" tur=\"{d['tur']}\">{d['mahkeme']} {d['esas']} | karşı taraf: {karsi.get(d['karsi'], '-')} | "
            f"son işlem {d['son_tarih'] or '-'}: {d['son'] or '-'} | sonraki işlem {d['sonraki_tarih'] or '-'}: {d['sonraki'] or '-'}"
            + (f" | alacak {tl(d['alacak_v'])}, tahsil edilen {tl(d['tahsil_v'] or 0)}, kalan {tl(d['kalan_v'])}" if d["alacak_v"] is not None else "")
            + (f" | {h} gündür işlem yok (hareketsiz)" if h is not None and h > hareketsiz else "") + "</dosya>")


def denetle(gruplar: list[dict], sonuc: dict, hareketsiz: int) -> list[str]:
    sorun = []
    for g in gruplar:
        a, t = g["kimlik"], sonuc.get(g["kimlik"])
        if not t:
            sorun.append(f"{a}: rapor yazılmadı")
            continue
        kendi = {d["dosya"]: d for d in g["dosyalar"]}
        diger = [d for x in gruplar if x is not g for d in x["dosyalar"]]
        yazilan = {x["dosya_no"]: x["durum"] for x in t["dosyalar"]}
        tum = " ".join([t["genel_ozet"]] + list(yazilan.values()) + t["muvekkilden_beklenenler"])
        eksik = [n for n in kendi if n not in yazilan]
        if eksik:
            sorun.append(f"{a}: dosya raporda yok: {', '.join(eksik)}")
        for n in yazilan:
            if n not in kendi:
                sorun.append(f"{a}: bu müvekkile ait olmayan dosya: {n}")
        sizan = sorted({x for d in diger for x in (d["dosya"], d["esas"]) if x and re.sub(r"\s*E\.?$", "", x) in tum})
        if sizan:
            sorun.append(f"{a}: başka müvekkilin dosya / esas numarası rapora karışmış: {', '.join(sizan)}")
        kodlar = set(re.findall(r"\{KT\d+\}", tum))
        if kodlar - set(g["karsi"].values()):
            sorun.append(f"{a}: bu müvekkilin dosyalarında olmayan karşı taraf kodu: {', '.join(sorted(kodlar - set(g['karsi'].values())))}")
        for n, metin in yazilan.items():
            d = kendi.get(n)
            if not d:
                continue
            fazla = sorted(sayilar(re.sub(r"\{[A-Z]+\d*\}", " ", metin)) - dosya_rakamlari(d))
            if fazla:
                sorun.append(f"{a}: {n}: dosya kaydında olmayan tarih / tutar: {', '.join(fazla)}")
            if d["sonraki_tarih_d"] and d["sonraki_tarih"] not in metin:
                sorun.append(f"{a}: {n}: sonraki işlem tarihi ({d['sonraki_tarih']}) yazılmamış")
        toplam = set().union(set(), *(dosya_rakamlari(d) for d in g["dosyalar"])) | {str(len(g["dosyalar"]))}
        toplam |= sayilar(tl(sum(d["kalan_v"] for d in g["dosyalar"] if d["kalan_v"] is not None)))
        fazla = sorted(sayilar(re.sub(r"\{[A-Z]+\d*\}", " ", t["genel_ozet"] + " " + " ".join(t["muvekkilden_beklenenler"]))) - toplam)
        if fazla:
            sorun.append(f"{a}: genel özette kayıtlarda olmayan sayı: {', '.join(fazla)}")
        if VAAT.search(tum):
            sorun.append(f"{a}: sonuç vaadi / kesinlik ifadesi ('{VAAT.search(tum).group(0)}'); yargılamanın sonucu önceden söylenemez")
        if "{MUVEKKIL}" not in t["genel_ozet"]:
            sorun.append(f"{a}: genel özette {{MUVEKKIL}} hitabı yok")
    return sorun


def uret(gruplar: list[dict], hareketsiz: int, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    sonuc, hedef = {}, gruplar
    if duzeltme is not None:  # düzeltmede yalnız sorunlu müvekkillerin raporu yeniden yazılır
        sorunlu = {s.split(":")[0] for s in duzeltme}
        sonuc.update({a: b for a, b in (onceki or {}).items() if a not in sorunlu})
        hedef = [g for g in gruplar if g["kimlik"] in sorunlu]
    for g in hedef:
        d = [s for s in (duzeltme or []) if s.split(":")[0] == g["kimlik"]]
        satir = "\n".join(dosya_satiri(x, g["karsi"], hareketsiz) for x in g["dosyalar"])
        mesaj = f"<muvekkil kod=\"{{MUVEKKIL}}\">\n{satir}\n</muvekkil>" + duzeltme_blogu(d, (onceki or {}).get(g["kimlik"]))
        sonuc[g["kimlik"]] = llm.json_iste(sistem, mesaj, SEMA)
    return sonuc


def doldur(metin: str, g: dict) -> str:
    metin = metin.replace("{MUVEKKIL}", g["ad"])
    for ad, kod in g["karsi"].items():
        metin = metin.replace(kod, ad)
    return metin


def rapor_metni(g: dict, t: dict, buro: str, bugun: date) -> str:
    d = {x["dosya"]: x for x in g["dosyalar"]}
    s = [f"{buro}", f"Tarih: {bugun:%d.%m.%Y}", "", f"DOSYA DURUM RAPORU — {g['ad']}", "", doldur(t["genel_ozet"], g), ""]
    for x in t["dosyalar"]:
        k = d.get(x["dosya_no"])
        if k:
            s += [f"■ {k['dosya']} · {k['mahkeme']} {k['esas']} ({k['tur']})", doldur(x["durum"], g), ""]
    if t["muvekkilden_beklenenler"]:
        s += ["Sizden beklediklerimiz:"] + [f"- {doldur(x, g)}" for x in t["muvekkilden_beklenenler"]] + [""]
    return "\n".join(s + ["Bu rapor bilgilendirme amaçlıdır; dosyaların sonucu hakkında bir taahhüt içermez. Sorularınız için büromuza ulaşabilirsiniz."])


def rapor_yaz(cikti: Path, gruplar, sonuc, kalan, tur, buro, bugun, hareketsiz) -> list[Path]:
    wb = Workbook()
    ws = wb.active
    ws.title = "Dosyalar"
    baslik(ws, ["Müvekkil", "Dosya No", "Tür", "Mahkeme / Daire", "Esas No", "Son İşlem", "Hareketsiz (gün)", "Sonraki İşlem", "Kalan Gün", "Kalan Alacak",
                "Rapordaki Açıklama"], (24, 11, 7, 30, 13, 40, 10, 34, 9, 15, 60))
    dosyalar = []
    for g in gruplar:
        t = sonuc.get(g["kimlik"]) or {"dosyalar": []}
        y = {x["dosya_no"]: x["durum"] for x in t["dosyalar"]}
        for d in g["dosyalar"]:
            ws.append([g["ad"], d["dosya"], d["tur"], d["mahkeme"], d["esas"], f"{d['son_tarih']} {d['son']}".strip(), d["hareketsiz_gun"],
                       f"{d['sonraki_tarih']} {d['sonraki']}".strip(), d["kalan_gun"], d["kalan_v"], doldur(y.get(d["dosya"], ""), g)])
            ws.cell(ws.max_row, 10).number_format = "#,##0.00"
            boya(ws, [11], AI_DOLGU)
            if d["hareketsiz_gun"] is not None and d["hareketsiz_gun"] > hareketsiz:
                boya(ws, [7], KIRMIZI)
            if d["kalan_gun"] is not None and 0 <= d["kalan_gun"] <= 14:
                boya(ws, [9], SARI)
        sor = [s for s in kalan if s.startswith(g["kimlik"] + ":")]
        if g["kimlik"] in sonuc and not sor:
            yol = cikti.with_name(f"{cikti.stem}_{re.sub(r'[^A-Za-z0-9]+', '_', katla(g['ad']))[:40]}.txt")
            yol.parent.mkdir(parents=True, exist_ok=True)
            yol.write_text(rapor_metni(g, sonuc[g["kimlik"]], buro, bugun), encoding="utf-8")
            dosyalar.append(yol)
    son_not(ws, "Önemli: Raporlar gönderilmedi. Avukat okuyup onaylamadan müvekkile iletmeyin; dosya kayıtlarını UYAP ile karşılaştırın.", kalan, tur)
    hizala(ws)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    return dosyalar


def calistir(girdi: Path, cikti: Path, buro: str, bugun: date | None = None, hareketsiz: int = 90, tur: int = 2, evet: bool = False) -> dict:
    bugun = bugun or date.today()
    dosyalar = dosyalari_oku(girdi, bugun)
    gruplar = muvekkiller(dosyalar)
    if not gruplar:
        raise llm.LLMHatasi(f"{girdi.name}: dosya yok.")
    hk = sum(1 for d in dosyalar if d["hareketsiz_gun"] is not None and d["hareketsiz_gun"] > hareketsiz)
    print(f"[OK] {len(dosyalar)} dosya · {len(gruplar)} müvekkil · {hk} hareketsiz dosya (>{hareketsiz} gün)")
    llm.onay_al(f"{len(dosyalar)} dosyanın mahkeme, esas no, işlem ve tutar bilgileri gönderilecek; müvekkil ve karşı taraf adları kodlanır. "
                "Avukatlık meslek kuralları ve müvekkil sırrı açısından uygunluğu değerlendirin.", evet)
    sonuc, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(gruplar, hareketsiz, d, o), lambda x: denetle(gruplar, x, hareketsiz), tur)
    yazilan = rapor_yaz(cikti, gruplar, sonuc, kalan, yapilan, buro, bugun, hareketsiz)
    return {"gruplar": gruplar, "sonuc": sonuc, "kalan": kalan, "tur": yapilan, "dosyalar": yazilan}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    p = argparse.ArgumentParser(description="Dava ve icra dosyalarının son durumunu müvekkile gönderilecek rapor biçiminde hazırlar (göndermez).")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "dosyalar.csv",
                   help="Dosya listesi (.xlsx/.csv): Dosya No, Müvekkil; Tür, Mahkeme / Daire, Esas No, Karşı Taraf, Son İşlem Tarihi, Son İşlem, "
                        "Sonraki İşlem Tarihi, Sonraki İşlem, Alacak Tutarı, Tahsil Edilen")
    p.add_argument("--buro", default="Örnek Hukuk Bürosu", help="Rapor başlığındaki büro adı")
    p.add_argument("--bugun", help="Rapor tarihi (GG.AA.YYYY; varsayılan: bugün)")
    p.add_argument("--hareketsiz", type=int, default=90, help="Bu kadar günden uzun süredir işlem görmeyen dosya hareketsiz sayılır (varsayılan 90)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "dosya_durum_raporu.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        bugun = tarih_coz(a.bugun) if a.bugun else None
        if a.bugun and not bugun:
            raise llm.LLMHatasi(f"--bugun '{a.bugun}' tarih olarak okunamadı (GG.AA.YYYY).")
        s = calistir(a.girdi, a.cikti, a.buro, bugun, a.hareketsiz, a.tur, a.evet)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(s['dosyalar'])}/{len(s['gruplar'])} müvekkil raporu hazır (gönderilmedi)")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
