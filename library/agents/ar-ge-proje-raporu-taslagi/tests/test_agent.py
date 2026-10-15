"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

V = KLASOR / "ornek_veri"
IP, BUT, NOT = V / "is_paketleri.csv", V / "butce.csv", V / "notlar.txt"
MESAJLAR = []


def ip(no, yap, son, sapma, sonraki, *d):
    return {"ip": no, "yapilanlar": yap, "sonuclar": son, "sapma_aciklamasi": sapma, "sonraki_donem": sonraki, "dayanak": list(d)}


def taslak(dogru=True) -> dict:
    t = {"donem_ozeti": "Dönemde sistem tasarımı tamamlanmış, sensör kartının ilk sürümü üretilmiş ve rulman arızası modelinde %91 doğruluk elde edilmiştir. "
                        "Sensör tedarikindeki gecikme İP2 ve İP4'ü etkilemiştir.",
         "is_paketleri": [
             ip("İP1", "Gereksinimler toplanmış, sistem mimarisi dokümanı hazırlanmıştır.", "Tasarım gözden geçirme toplantısı 28.05.2026 tarihinde yapılmıştır.", "",
                "-", "N1"),
             ip("İP2", "Sensör kartının ilk sürümü üretilmiş; veri toplama ve kablosuz iletim modülleri tamamlanmıştır.", "Gömülü yazılım modülleri çalışır durumdadır.",
                "İlerleme planlanan %80'e karşı %65'tir. İvmeölçer tedarik süresinin 6 haftadan 14 haftaya çıkması ve gürültü nedeniyle kart revizyonu gecikmeye yol açmıştır.",
                "Revize kart ve alternatif sensör uyumluluk testleri tamamlanacaktır.", "N2", "N3", "N4"),
             ip("İP3", "Laboratuvar düzeneğinde 3 pompadan veri toplanmış, rulman arızası sınıflandırma modeli eğitilmiştir.", "Doğrulama setinde %91 doğruluk elde edilmiştir.",
                "", "Kavitasyon arızası için ek veri toplanacaktır.", "N5", "N6"),
             ip("İP4", "Saha testleri için 2 müşteri sahası belirlenmiştir.", "-", "Prototip üretimi kart revizyonuna bağlı olarak 4 hafta ertelenmiştir.",
                "Prototip üretimine başlanacaktır.", "N7")],
         "butce_aciklamalari": [
             {"kalem": "Alet, teçhizat, yazılım", "aciklama": "Alternatif sensör ve test düzeneği için ek ölçüm cihazı alınmıştır.", "dayanak": ["N8"]},
             {"kalem": "Malzeme", "aciklama": "Malzeme alımları kart revizyonu nedeniyle sonraki döneme kaymıştır.", "dayanak": ["N9"]},
             {"kalem": "Seyahat", "aciklama": "Saha ziyaretleri çevrim içi yapıldığından harcama düşük kalmıştır.", "dayanak": ["N10"]}],
         "sorunlar_cozumler": ["MEMS ivmeölçer tedarik gecikmesi: alternatif sensör için uyumluluk testleri başlatılmıştır."],
         "ciktilar": []}
    if not dogru:
        t["is_paketleri"][1]["yapilanlar"] = "İP2 başarıyla tamamlanmıştır; sensör kartı 12 haftada üretilmiştir."
        t["is_paketleri"][1]["sapma_aciklamasi"] = ""
        del t["is_paketleri"][3]
        del t["butce_aciklamalari"][2]
        t["ciktilar"] = ["Sensör kartı için patent başvurusu yapılmıştır."]
    return t


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return taslak("<duzeltme>" in kullanici)


def veri():
    k, n = agent.notlari_oku(NOT)
    return agent.ip_oku(IP), agent.butce_oku(BUT, 10), k, n


class KodTesti(unittest.TestCase):
    def test_hesap(self):
        ipler, butce, k, n = veri()
        self.assertEqual([(x["no"], x["sapma"], x["durum"]) for x in ipler],
                         [("İP1", 0, "Tamamlandı"), ("İP2", -15, "Gecikmeli"), ("İP3", 5, "Planlandığı gibi"), ("İP4", -10, "Gecikmeli")])
        self.assertEqual([b["kalem"] for b in butce if b["aciklama_gerekli"]], ["Alet, teçhizat, yazılım", "Malzeme", "Seyahat"])
        self.assertEqual(round(butce[1]["oran"], 1), 22.6)
        self.assertEqual((k["program"], len(n), n[1]["konu"], n[7]["konu"]), ("TÜBİTAK TEYDEB 1501", 10, "İP2", "Bütçe"))

    def test_denetim(self):
        ipler, butce, k, n = veri()
        self.assertEqual(agent.denetle(taslak(), ipler, butce, n), [])
        s = agent.denetle(taslak(False), ipler, butce, n)
        for parca in ("İP2: notlarda / tabloda olmayan sayı: 12", "'tamamlandı'", "İP2: gecikmeli", "İP4: iş paketi raporda", "'Seyahat'", "patent"):
            self.assertTrue(any(parca in x for x in s), (parca, s))
        self.assertEqual(len(s), 6, s)


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "r.xlsx"
            s = agent.calistir(IP, BUT, NOT, cikti, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            self.assertNotIn("0000000", MESAJLAR[0])
            m = cikti.with_suffix(".txt").read_text(encoding="utf-8")
            self.assertIn("İP2 Sensör kartı ve gömülü yazılım geliştirme: planlanan %80, gerçekleşen %65 (Gecikmeli)", m)
            self.assertIn("Toplam: planlanan 2.740.000 TL, gerçekleşen 2.692.000 TL (gerçekleşme %98.2)", m)
            self.assertIn("Bu dönemde yok.", m)
            ws = load_workbook(cikti)["Bütçe"]
            self.assertEqual((ws["A3"].value, ws["E3"].value), ("Alet, teçhizat, yazılım", 22.6))

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(IP, BUT, NOT, Path(tmp) / "r.xlsx", tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 6)
            self.assertEqual(agent.main(["--evet", "--esik", "25", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--evet", "--notlar", str(Path(tmp) / "yok.txt")]), 1)


if __name__ == "__main__":
    unittest.main()
