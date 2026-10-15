"""
Ar-Ge Proje Raporu Taslağı — Workers / Workless AI Agent
Ar-Ge ve Ürün Geliştirme › Ar-Ge Mühendisi

1. Kod iş paketi tablosunu (planlanan / gerçekleşen ilerleme), bütçe tablosunu (planlanan / gerçekleşen harcama) ve proje
   notlarını okur. Notlar N1, N2 … diye numaralanır ("İP2: …" ile başlayan not o iş paketine bağlanır).
2. Kod hesaplar: her iş paketinin ilerleme sapması ve durumu (sapma −10 puan ve altı → gecikmeli), bütçe kalemlerinin
   sapması (± %10 üzeri → açıklama gerekli), toplam harcama gerçekleşme oranı.
3. Model dönem raporunun bölümlerini yazar: dönem özeti, iş paketi bölümleri (yapılanlar, sonuçlar, sapma açıklaması,
   sonraki dönem planı), bütçe sapma açıklamaları, sorunlar ve çözümler, proje çıktıları.
4. Kod denetler:
   - Her iş paketi ele alınmış mı; gecikmeli iş paketinde sapma açıklaması var mı
   - İş paketi bölümündeki sayılar o paketin notlarında / tablosunda ya da kod hesaplarında var mı; bölüm not numarası gösteriyor mu
   - %100'e ulaşmamış iş paketi için "tamamlandı" denmiş mi
   - Açıklama gereken her bütçe kalemi açıklanmış mı
   - Notlarda geçmeyen patent, yayın, ürün ya da ticarileşme iddiası var mı
   Sorun varsa modele geri gönderilir (en fazla --tur kez).
5. Kod rapor metnini proje künyesi, iş paketi ilerleme tablosu ve bütçe tablosuyla birlikte oluşturur.

Kullanım:
    python agent.py                                                         # örnek (kurgusal) kestirimci bakım projesi
    python agent.py --ip is_paketleri.xlsx --butce butce.xlsx --notlar notlar.docx
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
IP = {"no": ("İP No", "İş Paketi No", "No"), "ad": ("İş Paketi", "İş Paketi Adı", "Ad"), "bas": ("Planlanan Başlangıç", "Başlangıç"),
      "bit": ("Planlanan Bitiş", "Bitiş"), "plan": ("Planlanan İlerleme", "Planlanan %", "Planlanan"), "gercek": ("Gerçekleşen İlerleme", "Gerçekleşen %", "Gerçekleşen")}
BUTCE = {"kalem": ("Kalem", "Gider Kalemi"), "plan": ("Planlanan", "Bütçe"), "gercek": ("Gerçekleşen", "Harcama")}
KUNYE = {"ad": "Proje Adı", "no": "Proje No", "program": "Program", "donem": "Dönem"}
TAMAM = re.compile(r"\b(is paketi|ip ?\d+)\w*( basariyla)? (tamamlan(?!acak|masi)|bitiril|sonuclandiril)")
IDDIA = re.compile(r"\b(patent\w*|faydali model\w*|makale\w*|yayin\w*|bildiri\w*|ticarilest\w*|satis\w*|lisans\w*|tescil\w*)\b")

SEMA = {
    "type": "object",
    "properties": {
        "donem_ozeti": S,
        "is_paketleri": {"type": "array", "items": {"type": "object", "properties": {
            "ip": S, "yapilanlar": S, "sonuclar": S, "sapma_aciklamasi": S, "sonraki_donem": S, "dayanak": {"type": "array", "items": S}},
            "required": ["ip", "yapilanlar", "sonuclar", "sapma_aciklamasi", "sonraki_donem", "dayanak"], "additionalProperties": False}},
        "butce_aciklamalari": {"type": "array", "items": {"type": "object", "properties": {"kalem": S, "aciklama": S, "dayanak": {"type": "array", "items": S}},
                                                          "required": ["kalem", "aciklama", "dayanak"], "additionalProperties": False}},
        "sorunlar_cozumler": {"type": "array", "items": S},
        "ciktilar": {"type": "array", "items": S},
    },
    "required": ["donem_ozeti", "is_paketleri", "butce_aciklamalari", "sorunlar_cozumler", "ciktilar"],
    "additionalProperties": False,
}


def sayi(x) -> float:
    if isinstance(x, (int, float)):
        return float(x)
    try:
        return float(sayi_norm(str(x).replace("%", "").replace("TL", "").strip() or "0"))
    except ValueError:
        raise llm.LLMHatasi(f"Sayı okunamadı: '{x}'")


def para(v: float) -> str:
    return f"{v:,.0f}".replace(",", ".")


def ip_oku(yol: Path) -> list[dict]:
    out = []
    for r in kayitlar(yol, IP, ("no", "plan", "gercek")):
        if not r["no"]:
            continue
        x = {"no": str(r["no"]).strip(), "ad": str(r["ad"]), "bas": tarih_coz(r["bas"]), "bit": tarih_coz(r["bit"]), "plan": sayi(r["plan"]), "gercek": sayi(r["gercek"])}
        x["sapma"] = x["gercek"] - x["plan"]
        x["durum"] = "Tamamlandı" if x["gercek"] >= 100 else ("Gecikmeli" if x["sapma"] <= -10 else ("Planın önünde" if x["sapma"] >= 10 else "Planlandığı gibi"))
        out.append(x)
    if not out:
        raise llm.LLMHatasi(f"{yol.name}: iş paketi yok.")
    return out


def butce_oku(yol: Path, esik: float) -> list[dict]:
    out = []
    for r in kayitlar(yol, BUTCE, ("kalem", "plan", "gercek")):
        if not r["kalem"]:
            continue
        x = {"kalem": str(r["kalem"]), "plan": sayi(r["plan"]), "gercek": sayi(r["gercek"])}
        x["fark"] = x["gercek"] - x["plan"]
        x["oran"] = x["fark"] / x["plan"] * 100 if x["plan"] else 0
        x["aciklama_gerekli"] = abs(x["oran"]) > esik
        out.append(x)
    return out


def notlari_oku(yol: Path) -> tuple[dict, list[dict]]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: proje notları bulunamadı.")
    k = {a: "" for a in KUNYE}
    notlar, bolum = [], False
    for s in belge.metin_oku(yol).splitlines():
        if not s.strip():
            continue
        a, _, d = s.partition(":")
        anahtar = next((x for x, v in KUNYE.items() if katla(v) == katla(a)), None)
        if anahtar and not s.lstrip().startswith(("-", "•", "*")):
            k[anahtar] = d.strip()
        elif katla(s) == "notlar":
            bolum = True
        elif bolum:
            m = s.strip().lstrip("-•* ").strip()
            etiket = re.match(r"([^:]{1,15}):\s*(.+)", m)
            notlar.append({"id": f"N{len(notlar) + 1}", "konu": etiket.group(1).strip() if etiket else "", "metin": etiket.group(2) if etiket else m})
    if not notlar:
        raise llm.LLMHatasi(f"{yol.name}: 'Notlar' bölümü gerekli ('- İP2: …' satırları).")
    return k, notlar


def ip_satiri(x: dict) -> str:
    return (f"{x['no']} | {x['ad']} | plan {x['bas']:%d.%m.%Y}-{x['bit']:%d.%m.%Y} | planlanan ilerleme %{x['plan']:.0f} | gerçekleşen %{x['gercek']:.0f} | "
            f"sapma {x['sapma']:+.0f} puan | {x['durum']}") if x["bas"] and x["bit"] else (
            f"{x['no']} | {x['ad']} | planlanan %{x['plan']:.0f} | gerçekleşen %{x['gercek']:.0f} | sapma {x['sapma']:+.0f} puan | {x['durum']}")


def butce_satiri(b: dict) -> str:
    return f"{b['kalem']} | planlanan {para(b['plan'])} TL | gerçekleşen {para(b['gercek'])} TL | fark {para(b['fark'])} TL (%{b['oran']:+.1f})"


def denetle(t: dict, ipler: list[dict], butce: list[dict], notlar: list[dict]) -> list[str]:
    sorun = []
    ip = {x["no"]: x for x in ipler}
    n = {x["id"]: x for x in notlar}
    tum_not = katla(" ".join(x["metin"] for x in notlar))
    genel = " ".join(ip_satiri(x) for x in ipler) + " " + " ".join(butce_satiri(b) for b in butce)
    gorulen = set()
    for b in t["is_paketleri"]:
        if b["ip"] not in ip:
            sorun.append(f"{b['ip']}: iş paketi tabloda yok")
            continue
        gorulen.add(b["ip"])
        x = ip[b["ip"]]
        gec = [d for d in b["dayanak"] if d not in n]
        if not b["dayanak"] or gec:
            sorun.append(f"{b['ip']}: not dayanağı yok ya da geçersiz {gec}")
        metin = " ".join([b["yapilanlar"], b["sonuclar"], b["sapma_aciklamasi"], b["sonraki_donem"]])
        kaynak = ip_satiri(x) + " " + " ".join(n[d]["metin"] for d in b["dayanak"] if d in n)
        fazla = sorted(sayilar(metin) - sayilar(kaynak))
        if fazla:
            sorun.append(f"{b['ip']}: notlarda / tabloda olmayan sayı: {', '.join(fazla)}")
        if x["gercek"] < 100 and TAMAM.search(katla(b["yapilanlar"] + " " + b["sonuclar"])):
            sorun.append(f"{b['ip']}: ilerleme %{x['gercek']:.0f} iken iş paketi 'tamamlandı' denmiş")
        if x["durum"] == "Gecikmeli" and len(b["sapma_aciklamasi"].strip()) < 20:
            sorun.append(f"{b['ip']}: gecikmeli ({x['sapma']:+.0f} puan) ama sapma açıklaması yok")
    for k in ip:
        if k not in gorulen:
            sorun.append(f"{k}: iş paketi raporda ele alınmamış")
    aciklanan = {katla(a["kalem"]) for a in t["butce_aciklamalari"]}
    for b in butce:
        if b["aciklama_gerekli"] and katla(b["kalem"]) not in aciklanan:
            sorun.append(f"Bütçe: '{b['kalem']}' kaleminde %{b['oran']:+.1f} sapma var, açıklanmamış")
    for a in t["butce_aciklamalari"]:
        gec = [d for d in a["dayanak"] if d not in n]
        if not a["dayanak"] or gec:
            sorun.append(f"Bütçe / {a['kalem']}: not dayanağı yok ya da geçersiz {gec}")
        kaynak = genel + " " + " ".join(n[d]["metin"] for d in a["dayanak"] if d in n)
        fazla = sorted(sayilar(a["aciklama"]) - sayilar(kaynak))
        if fazla:
            sorun.append(f"Bütçe / {a['kalem']}: tablolarda / notlarda olmayan sayı: {', '.join(fazla)}")
    serbest = " ".join([t["donem_ozeti"]] + t["sorunlar_cozumler"] + t["ciktilar"] + [b["sonuclar"] for b in t["is_paketleri"]])
    for m in IDDIA.finditer(katla(serbest)):
        if m.group(0)[:6] not in tum_not:
            sorun.append(f"Çıktı / iddia: '{m.group(0)}' notlarda geçmiyor (patent, yayın, ticarileşme uydurulamaz)")
    fazla = sorted(sayilar(t["donem_ozeti"] + " " + " ".join(t["sorunlar_cozumler"] + t["ciktilar"])) - sayilar(genel + " " + " ".join(x["metin"] for x in notlar)))
    if fazla:
        sorun.append(f"Özet / sorunlar / çıktılar: kaynaklarda olmayan sayı: {', '.join(fazla)}")
    return list(dict.fromkeys(sorun))


def uret(k, ipler, butce, notlar, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    mesaj = (f"<proje>{k['ad']} · {k['program'] or '-'} · dönem {k['donem'] or '-'}</proje>\n<is_paketleri>\n" + "\n".join(ip_satiri(x) for x in ipler) +
             "\n</is_paketleri>\n<butce>\n" + "\n".join(butce_satiri(b) + (" | AÇIKLAMA GEREKLİ" if b["aciklama_gerekli"] else "") for b in butce) +
             "\n</butce>\n<notlar>\n" + "\n".join(f"{x['id']} | {x['konu']} | {x['metin']}" for x in notlar) + "\n</notlar>" + duzeltme_blogu(duzeltme, onceki))
    return llm.json_iste(sistem, mesaj, SEMA)


def rapor_metni(t, k, ipler, butce) -> str:
    tp, tg = sum(b["plan"] for b in butce), sum(b["gercek"] for b in butce)
    s = ["AR-GE PROJESİ DÖNEM RAPORU (TASLAK)", "", f"Proje adı: {k['ad']}", f"Proje no: {k['no']}", f"Program: {k['program']}", f"Rapor dönemi: {k['donem']}",
         "", "1. DÖNEM ÖZETİ", t["donem_ozeti"], "", "2. İŞ PAKETLERİ", ""]
    s += [f"  {x['no']} {x['ad']}: planlanan %{x['plan']:.0f}, gerçekleşen %{x['gercek']:.0f} ({x['durum']})" for x in ipler]
    ip = {x["no"]: x for x in ipler}
    for b in t["is_paketleri"]:
        s += ["", f"{b['ip']} {ip.get(b['ip'], {}).get('ad', '')}", f"Yapılan çalışmalar: {b['yapilanlar']}", f"Elde edilen sonuçlar: {b['sonuclar']}"]
        if b["sapma_aciklamasi"].strip():
            s.append(f"Planlanandan sapma ve nedenleri: {b['sapma_aciklamasi']}")
        s.append(f"Sonraki dönem planı: {b['sonraki_donem']}")
    s += ["", "3. HARCAMALAR"] + [f"  {b['kalem']}: planlanan {para(b['plan'])} TL, gerçekleşen {para(b['gercek'])} TL (%{b['oran']:+.1f})" for b in butce]
    if tp:
        s.append(f"  Toplam: planlanan {para(tp)} TL, gerçekleşen {para(tg)} TL (gerçekleşme %{tg / tp * 100:.1f})")
    s += [f"  {a['kalem']}: {a['aciklama']}" for a in t["butce_aciklamalari"]]
    s += ["", "4. KARŞILAŞILAN SORUNLAR VE ÇÖZÜMLER"] + [f"  - {x}" for x in t["sorunlar_cozumler"]]
    s += ["", "5. PROJE ÇIKTILARI"] + ([f"  - {x}" for x in t["ciktilar"]] or ["  Bu dönemde yok."])
    return "\n".join(s)


def rapor_yaz(cikti: Path, t, k, ipler, butce, notlar, kalan, tur, metin) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "İş Paketleri"
    baslik(ws, ["İP", "Ad", "Planlanan %", "Gerçekleşen %", "Sapma (puan)", "Durum", "Yapılanlar", "Sonuçlar", "Sapma Açıklaması", "Sonraki Dönem", "Dayanak"],
           (6, 30, 10, 10, 10, 14, 50, 40, 40, 40, 12))
    t_ip = {b["ip"]: b for b in t["is_paketleri"]}
    for x in ipler:
        b = t_ip.get(x["no"], {})
        ws.append([x["no"], x["ad"], x["plan"], x["gercek"], x["sapma"], x["durum"], b.get("yapilanlar", ""), b.get("sonuclar", ""),
                   b.get("sapma_aciklamasi", ""), b.get("sonraki_donem", ""), ", ".join(b.get("dayanak", []))])
        boya(ws, (6,), {"Gecikmeli": KIRMIZI, "Tamamlandı": YESIL}.get(x["durum"], SARI if x["durum"] == "Planın önünde" else YESIL))
        boya(ws, (7, 8, 9, 10), AI_DOLGU)
    son_not(ws, "Önemli: Taslaktır. Teknik içerik proje yürütücüsü, harcamalar mali sorumlu tarafından doğrulanmalı; rapor kurumun (TÜBİTAK, "
                "Sanayi ve Teknoloji Bakanlığı vb.) güncel formuna aktarılmalıdır.", kalan, tur)
    hizala(ws)
    w = wb.create_sheet("Bütçe")
    baslik(w, ["Kalem", "Planlanan", "Gerçekleşen", "Fark", "Fark %", "Açıklama"], (26, 14, 14, 14, 10, 70))
    acik = {katla(a["kalem"]): a["aciklama"] for a in t["butce_aciklamalari"]}
    for b in butce:
        w.append([b["kalem"], b["plan"], b["gercek"], b["fark"], round(b["oran"], 1), acik.get(katla(b["kalem"]), "")])
        for c in (2, 3, 4):
            w.cell(w.max_row, c).number_format = "#,##0"
        if b["aciklama_gerekli"]:
            boya(w, (5,), SARI)
    hizala(w)
    n = wb.create_sheet("Notlar")
    baslik(n, ["No", "Konu", "Not"], (6, 12, 110))
    for x in notlar:
        n.append([x["id"], x["konu"], x["metin"]])
    hizala(n)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    cikti.with_suffix(".txt").write_text(metin, encoding="utf-8")


def calistir(ip_yol: Path, butce_yol: Path, not_yol: Path, cikti: Path, esik: float = 10, tur: int = 2, evet: bool = False) -> dict:
    ipler = ip_oku(ip_yol)
    butce = butce_oku(butce_yol, esik)
    k, notlar = notlari_oku(not_yol)
    print(f"[OK] {k['ad'] or 'Proje'} · {len(ipler)} iş paketi · {len(butce)} bütçe kalemi · {len(notlar)} not")
    for x in ipler:
        if x["durum"] == "Gecikmeli":
            print(f"[!] {x['no']} gecikmeli ({x['sapma']:+.0f} puan)")
    llm.onay_al(f"Proje adı, iş paketi ve bütçe tabloları, {len(notlar)} not gönderilecek (proje no gönderilmez). Gizli teknik bilgi içeriyorsa dikkat edin.", evet)
    t, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(k, ipler, butce, notlar, d, o), lambda x: denetle(x, ipler, butce, notlar), tur)
    metin = rapor_metni(t, k, ipler, butce)
    rapor_yaz(cikti, t, k, ipler, butce, notlar, kalan, yapilan, metin)
    return {"ipler": ipler, "butce": butce, "taslak": t, "kalan": kalan, "tur": yapilan, "metin": metin}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    p = argparse.ArgumentParser(description="Proje notları ve sonuçlarından Ar-Ge dönem raporu taslağı hazırlar.")
    v = BURASI / "ornek_veri"
    p.add_argument("--ip", type=Path, default=v / "is_paketleri.csv", help="İş paketleri (CSV/Excel): İP No, İş Paketi, Planlanan Başlangıç / Bitiş, Planlanan İlerleme, Gerçekleşen İlerleme")
    p.add_argument("--butce", type=Path, default=v / "butce.csv", help="Bütçe (CSV/Excel): Kalem, Planlanan, Gerçekleşen")
    p.add_argument("--notlar", type=Path, default=v / "notlar.txt", help="Notlar (.txt/.docx): Proje Adı, Proje No, Program, Dönem satırları; 'Notlar' bölümü ('- İP2: …')")
    p.add_argument("--esik", type=float, default=10, help="Açıklama gerektiren bütçe sapması, %% (varsayılan 10)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "arge_donem_raporu.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.ip, a.butce, a.notlar, a.cikti, a.esik, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.with_suffix('.txt').resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
