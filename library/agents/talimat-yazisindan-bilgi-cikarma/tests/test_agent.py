"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import os
import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
GIDA = ORNEK / "talimat_ornek_gida.txt"
TEKSTIL = ORNEK / "talimat_kurgu_tekstil.txt"
MESAJLAR = []


def tal(**kw):
    t = {"referans": "", "tarih": "", "gonderen_unvan": "", "alici_unvan": "", "tutar": "", "tutar_yazi": "", "aciklama": "", "gonderen_iban": "",
         "alici_iban": "", "tur": "Belirtilmemiş", "para_birimi": "TL", "imzalayanlar": []}
    return t | kw


IMZA = ["Kemal Örnek", "Selin Örnek"]
GIDA_DOGRU = {"talimatlar": [
    tal(referans="OG-2026/412", tarih="14.10.2026", gonderen_unvan="Örnek Gıda San. ve Tic. A.Ş.", alici_unvan="Kurgu Lojistik A.Ş.", tutar="125.000,50",
        tutar_yazi="Yalnız Yüzyirmibeşbin Türk Lirası Elli Kuruş", aciklama="Fatura no 2026-118 ödemesi", gonderen_iban="[IBAN-1]", alici_iban="[IBAN-2]",
        tur="EFT", imzalayanlar=IMZA),
    tal(referans="OG-2026/412", tarih="14.10.2026", gonderen_unvan="Örnek Gıda San. ve Tic. A.Ş.", alici_unvan="Kurgu Nakliyat Ltd. Şti.", tutar="48.500,00",
        tutar_yazi="Yalnız Kırksekizbin Türk Lirası", aciklama="Ekim ayı nakliye bedeli", gonderen_iban="[IBAN-1]", alici_iban="[IBAN-3]", tur="EFT",
        imzalayanlar=IMZA)], "notlar": []}
TEKSTIL_DOGRU = {"talimatlar": [
    tal(tarih="15.10.2026", gonderen_unvan="Kurgu Tekstil Ltd. Şti.", alici_unvan="Ayşe Deneme", tutar="15.750,00", tutar_yazi="onbeşbinyediyüzelli Türk Lirası",
        aciklama="Kira bedeli Ekim 2026", gonderen_iban="[IBAN-1]", alici_iban="[IBAN-2]", tur="Havale", imzalayanlar=["Can Kurgu"])],
    "notlar": ["Talimat e-posta ile iletilmiş."]}


def gida_hatali() -> dict:
    c = copy.deepcopy(GIDA_DOGRU)
    c["talimatlar"][0]["alici_unvan"] = "Kurgu Lojistik AŞ"          # belgedeki yazım değil
    c["talimatlar"][1]["alici_iban"] = "TR330099900001000000400001"  # yer tutucu yerine IBAN uydurmuş
    return c


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    if "Kurgu Tekstil" in kullanici:
        c = copy.deepcopy(TEKSTIL_DOGRU)
        if "[KİŞİ-1]" in kullanici:                   # --gizle: model adı yalnız takma adıyla görür
            c["talimatlar"][0]["alici_unvan"] = "[KİŞİ-1]"
        return c
    return copy.deepcopy(GIDA_DOGRU) if "<duzeltme>" in kullanici else gida_hatali()


class KodTesti(unittest.TestCase):
    def test_iban(self):
        self.assertEqual(agent.iban_kontrol("TR480099900001000000100001"), "")
        self.assertIn("mod 97", agent.iban_kontrol("TR380099900001000000300010"))
        self.assertIn("rezerv", agent.iban_kontrol("TR480009990000100000010001"))
        self.assertIn("26 karakter", agent.iban_kontrol("TR4800999000010000001000"))
        m, h = agent.iban_maskele("A TR48 0099 9000 0100 0000 1000 01 B tr040088800005500000000123 C TR480099900001000000100001 D TR12 3456")
        self.assertEqual(m, "A [IBAN-1] B [IBAN-2] C [IBAN-1] D TR12 3456")
        self.assertEqual(h, {"[IBAN-1]": "TR480099900001000000100001", "[IBAN-2]": "TR040088800005500000000123"})

    def test_yazidan_tutar(self):
        self.assertEqual(agent.yazidan_tutar("Yalnız Yüzyirmibeşbin Türk Lirası Elli Kuruş"), D("125000.50"))
        self.assertEqual(agent.yazidan_tutar("bir milyon iki yüz bin TL"), D("1200000"))
        self.assertEqual(agent.yazidan_tutar("YALNIZ KIRKSEKİZBİN TÜRK LİRASI"), D("48000"))
        self.assertEqual(agent.yazidan_tutar("48.500,00 TL"), D("48500.00"))
        self.assertIsNone(agent.yazidan_tutar("bir şey"))
        self.assertEqual((agent.para("125.000,50"), agent.para("125,000.50"), agent.para("48.500"), agent.para("abc")), (D("125000.50"), D("125000.50"), D("48500"), None))

    def test_maskeleme_ve_geri_acma(self):
        k = agent.takma_adlar(["Ayşe Deneme", "", "Ayşe Deneme"])
        m, h = agent.maskele(TEKSTIL.read_text(encoding="utf-8"), k)
        self.assertEqual(k, {"Ayşe Deneme": "[KİŞİ-1]"})
        for x in ("TR43", "TR38", "Ayşe Deneme", "can.kurgu@", "0532"):
            self.assertNotIn(x, m)
        self.assertIn("[KİŞİ-1] adına\n[IBAN-2]", m)
        self.assertEqual(agent.geri_ac("[KİŞİ-1] kira", k), "Ayşe Deneme kira")

    def test_denetim(self):
        m, h = agent.maskele(GIDA.read_text(encoding="utf-8"), {})
        self.assertEqual(agent.denetle(GIDA_DOGRU, m, h), [])
        s = agent.denetle(gida_hatali(), m, h)
        self.assertEqual(len(s), 2)
        self.assertTrue(s[0].startswith("Talimat 1: Alıcı unvan") and s[1].startswith("Talimat 2: Alıcı IBAN"))
        c = copy.deepcopy(GIDA_DOGRU)
        c["talimatlar"][0].update(alici_iban="[IBAN-1]", tutar="", imzalayanlar=["Murat Kurgu"])
        c["talimatlar"][1]["tutar"] = "kırk bin"
        s = agent.denetle(c, m, h)
        self.assertEqual(len(s), 5)                        # aynı IBAN, boş tutar, imzalayan; 2. talimatta belgede olmayan + okunamayan tutar
        self.assertTrue(any("aynı" in x for x in s) and any("rakamla tutar boş" in x for x in s) and any("Murat Kurgu" in x for x in s))

    def test_bulgular(self):
        t = GIDA_DOGRU["talimatlar"][0] | {"gonderen_iban": "TR480099900001000000100001", "alici_iban": "TR040088800005500000000123"}
        self.assertEqual(agent.bulgular(t, "mektup"), [])
        t2 = t | {"alici_iban": "", "gonderen_iban": "TR380099900001000000300010", "tutar_yazi": "", "imzalayanlar": [], "tur": "Belirtilmemiş",
                  "para_birimi": "Belirtilmemiş", "tarih": "Ekim ortası"}
        self.assertEqual([o for o, _ in agent.bulgular(t2, "mektup")], ["Hata", "Hata", "Dikkat", "Hata", "Dikkat", "Dikkat", "Dikkat"])
        t3 = t | {"tutar": "48.500,00", "tutar_yazi": "Kırksekizbin TL"}
        self.assertEqual(agent.bulgular(t3, "mektup"), [("Hata", "Rakamla tutar 48.500,00, yazıyla tutar 48.000,00: uyuşmuyor")])
        self.assertIn("E-postayla", agent.bulgular(t, "Kimden: a\nKime: b\nKonu: c")[0][1])


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ornek_ve_duzeltme(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            s = agent.calistir(ORNEK, cikti, evet=True)
            self.assertEqual(len(MESAJLAR), 3)               # tekstil 1 istek, gıda 1 + 1 düzeltme
            self.assertIn("Alıcı IBAN 'TR330099900001000000400001'", MESAJLAR[2].split("<duzeltme>")[1])
            for m in MESAJLAR:
                self.assertNotRegex(m.split("<duzeltme>")[0], r"TR\d{2}\s?\d{4}")   # IBAN sağlayıcıya gitmez
            self.assertEqual([(t["no"], t["durum"], t["alici_iban"]) for t in s["talimatlar"]],
                             [("talimat_kurgu_tekstil-1", "Hatalı", "TR380099900001000000300010"), ("OG-2026/412/1", "Uygun", "TR040088800005500000000123"),
                              ("OG-2026/412/2", "Hatalı", "TR330099900001000000400001")])
            self.assertEqual([d["tur"] for d in s["dosyalar"]], [0, 1])
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Talimatlar", "Bulgular", "Kontrol"])
            ws = wb["Talimatlar"]
            self.assertEqual([c.value for c in ws[1]][:13], agent.KOD_BLOGU_SUTUNLARI)
            self.assertEqual([ws.cell(3, j).value for j in (1, 6, 8, 9, 13, 14)],
                             ["OG-2026/412/1", "TR480099900001000000100001", "TR040088800005500000000123", 125000.5, "Kemal Örnek, Selin Örnek", "Uygun"])
            self.assertIn('"no": "OG-2026/412/2"', s["json"].read_text(encoding="utf-8"))

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_gizle_duzeltme_kapali_ve_cli(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(GIDA, Path(tmp) / "t.xlsx", tur=0, evet=True)
            self.assertEqual(len(s["dosyalar"][0]["kalan"]), 2)
            self.assertEqual([t["durum"] for t in s["talimatlar"]], ["Kontrol", "Hatalı"])
            self.assertEqual(s["talimatlar"][1]["alici_iban"], "")                    # uydurma IBAN kullanılmaz
            n = len(MESAJLAR)
            s = agent.calistir(TEKSTIL, Path(tmp) / "t.xlsx", gizle=["Ayşe Deneme"], evet=True)
            self.assertEqual((len(MESAJLAR) - n, s["dosyalar"][0]["kalan"], s["talimatlar"][0]["alici_unvan"]), (1, [], "Ayşe Deneme"))
            self.assertNotIn("Ayşe", MESAJLAR[-1])
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--girdi", str(ORNEK / "yok.txt"), "--evet"]), 1)
            bos = Path(tmp) / "bos"
            bos.mkdir()
            self.assertEqual(agent.main(["--girdi", str(bos), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
