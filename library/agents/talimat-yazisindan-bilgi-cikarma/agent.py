"""
Talimat Yazısından Bilgi Çıkarma — Gerisi Bende AI Agent
Şube Bankacılığı › Gişe Yetkilisi

1. Kod, talimat belgesindeki (mektup, faks, e-posta; .pdf/.docx/.txt) TR IBAN'ları bulur ve [IBAN-1], [IBAN-2] … yer
   tutucularıyla değiştirir; IBAN'lar sağlayıcıya gönderilmez. Telefon, e-posta, TCKN ve --gizle ile verilen adlar da maskelenir.
2. Model her ödeme için gönderen, alıcı, IBAN yer tutucusu, tutar (rakam ve yazı), para birimi, işlem türü, açıklama,
   tarih ve imzalayanları çıkarır.
3. Kod çıkarımı denetler: her metin belgede aynen geçiyor mu, IBAN alanında belgede olan bir yer tutucu mu var, rakamla tutar
   okunabiliyor mu, gönderen ve alıcı IBAN'ı aynı mı. Sorun varsa modele geri gönderilir (en fazla --tur kez).
4. Kod gerçek IBAN'ları geri koyar ve belgeyle ilgili bulguları ekler: IBAN kontrol hanesi (ISO 13616 mod 97), rakam-yazı
   tutar uyumu, eksik alıcı IBAN'ı / imza / yazıyla tutar, e-postayla gelen talimat.
Çıktının ilk sayfası "Müşteri Talimat Kontrolü" kod bloğunun girdi biçimindedir (Müşteri No sütununu siz doldurursunuz).

Kullanım:
    python agent.py                                   # örnek talimatlarla dener
    python agent.py --girdi ./talimatlar --gizle "Ayşe Deneme"
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
RENK = {"Hata": "FDE2E1", "Dikkat": "FFF4CE", "Bilgi": "E8F0FE", "Hatalı": "FDE2E1", "Kontrol": "FFF4CE", "Uygun": "E3F5E1"}
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def para(x) -> Decimal | None:
    """'125.000,50' · '125,000.50' · '48500' → Decimal."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1 or (s.count(".") == 1 and len(s.split(".")[1]) == 3):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tarih_coz(x) -> date | None:
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


# ----------------------------------------------------------------------------
# IBAN ve yazıyla tutar
# ----------------------------------------------------------------------------

IBAN_DESEN = re.compile(r"\bTR\s?\d{2}(?:\s?[0-9]{4}){5}\s?\d{2}\b", re.I)


def iban_temizle(x) -> str:
    return re.sub(r"[\s-]", "", str(x or "")).upper()


def iban_kontrol(iban: str) -> str:
    """Boş metin = geçerli. ISO 13616 mod 97; TR IBAN 26 karakter, 10. karakter (rezerv) 0."""
    if not re.fullmatch(r"[A-Z]{2}\d{2}[A-Z0-9]{10,30}", iban):
        return "IBAN biçimi hatalı"
    if iban.startswith("TR"):
        if len(iban) != 26:
            return f"TR IBAN 26 karakter olmalı ({len(iban)})"
        if iban[9] != "0":
            return "TR IBAN'ın 10. karakteri (rezerv alan) 0 olmalı"
    if int("".join(str(int(c, 36)) for c in iban[4:] + iban[:4])) % 97 != 1:
        return "IBAN kontrol hanesi (mod 97) tutmuyor — yazım hatası olabilir"
    return ""


def iban_maskele(metin: str) -> tuple[str, dict[str, str]]:
    """TR IBAN'ları geçtiği sırayla [IBAN-n] yapar; aynı IBAN aynı yer tutucuyu alır. Dönen harita: yer tutucu → IBAN."""
    harita: dict[str, str] = {}
    ters: dict[str, str] = {}

    def degistir(m):
        iban = iban_temizle(m.group(0))
        if iban not in ters:
            ters[iban] = f"[IBAN-{len(ters) + 1}]"
            harita[ters[iban]] = iban
        return ters[iban]
    return IBAN_DESEN.sub(degistir, metin), harita


SAYI_SOZ = {"sifir": 0, "bir": 1, "iki": 2, "uc": 3, "dort": 4, "bes": 5, "alti": 6, "yedi": 7, "sekiz": 8, "dokuz": 9, "on": 10, "yirmi": 20,
            "otuz": 30, "kirk": 40, "elli": 50, "altmis": 60, "yetmis": 70, "seksen": 80, "doksan": 90, "yuz": 100, "bin": 1000,
            "milyon": 10 ** 6, "milyar": 10 ** 9}
_SOZ_SIRA = sorted(SAYI_SOZ, key=len, reverse=True)
ANA_BIRIM = {"tl", "try", "lira", "lirasi", "dolar", "usd", "avro", "euro", "eur", "sterlin", "gbp", "frank", "chf"}
ALT_BIRIM = {"kurus", "krs", "kr", "sent", "cent"}


def sozden_sayi(metin: str) -> int | None:
    """Bitişik veya ayrı yazılmış Türkçe sayı ('yüzyirmibeşbin', 'bir milyon iki yüz bin') → tam sayı."""
    s = re.sub(r"\s+", "", metin)
    if not s:
        return None
    toplam, simdiki, i = 0, 0, 0
    while i < len(s):
        soz = next((x for x in _SOZ_SIRA if s.startswith(x, i)), None)
        if soz is None:
            return None
        v = SAYI_SOZ[soz]
        i += len(soz)
        if v == 100:
            simdiki = (simdiki or 1) * 100
        elif v >= 1000:
            toplam += (simdiki or 1) * v
            simdiki = 0
        else:
            simdiki += v
    return toplam + simdiki


def yazidan_tutar(yazi) -> Decimal | None:
    """'Yalnız Yüzyirmibeşbin Türk Lirası Elli Kuruş' → 125000.50. Okunamazsa None."""
    s = str(yazi or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()
    if not s.strip():
        return None
    if re.search(r"\d", s):
        return para(re.sub(r"[^\d,.]", "", s))
    kelimeler = [k for k in re.split(r"[^a-z]+", s) if k and k not in {"yalniz", "yalnizca", "turk"}]
    ana, alt, bolum = [], [], 0
    for k in kelimeler:
        if k in ANA_BIRIM:
            bolum = 1
        elif k in ALT_BIRIM:
            bolum = 2
        elif bolum == 0:
            ana.append(k)
        elif bolum == 1:
            alt.append(k)
        else:
            return None
    a = sozden_sayi("".join(ana)) if ana else 0
    b = sozden_sayi("".join(alt)) if alt else 0
    if a is None or b is None or (not ana and not alt) or b >= 100:
        return None
    return Decimal(a) + Decimal(b) / 100


def tl(x: Decimal) -> str:
    return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

METIN_ALANLARI = ["referans", "tarih", "gonderen_unvan", "alici_unvan", "tutar", "tutar_yazi", "aciklama"]
ETIKET = {"referans": "Referans", "tarih": "Tarih", "gonderen_unvan": "Gönderen unvan", "alici_unvan": "Alıcı unvan", "tutar": "Tutar",
          "tutar_yazi": "Tutar (yazı)", "aciklama": "Açıklama", "gonderen_iban": "Gönderen IBAN", "alici_iban": "Alıcı IBAN"}
TALIMAT = {"type": "object",
           "properties": {**{a: S for a in METIN_ALANLARI}, "gonderen_iban": S, "alici_iban": S,
                          "tur": {"type": "string", "enum": ["EFT", "FAST", "Havale", "Virman", "Belirtilmemiş"]},
                          "para_birimi": {"type": "string", "enum": ["TL", "USD", "EUR", "GBP", "CHF", "Diğer", "Belirtilmemiş"]},
                          "imzalayanlar": {"type": "array", "items": S}},
           "required": [*METIN_ALANLARI, "gonderen_iban", "alici_iban", "tur", "para_birimi", "imzalayanlar"], "additionalProperties": False}
SEMA = {"type": "object", "properties": {"talimatlar": {"type": "array", "items": TALIMAT}, "notlar": {"type": "array", "items": S}},
        "required": ["talimatlar", "notlar"], "additionalProperties": False}


def takma_adlar(terimler: list[str]) -> dict[str, str]:
    return {ad: f"[KİŞİ-{i}]" for i, ad in enumerate(dict.fromkeys(t.strip() for t in terimler if t.strip()), 1)}


def maskele(metin: str, kisiler: dict[str, str]) -> tuple[str, dict[str, str]]:
    """Önce IBAN'lar ([IBAN-n]), sonra --gizle adları ([KİŞİ-n]), en son ortak maskeler (telefon, e-posta, TCKN)."""
    metin, ibanlar = iban_maskele(metin)
    for gercek, takma in sorted(kisiler.items(), key=lambda x: -len(x[0])):
        metin = re.sub(re.escape(gercek), takma, metin, flags=re.I)
    return llm.maskele(metin), ibanlar


def geri_ac(s: str, kisiler: dict[str, str]) -> str:
    for gercek, takma in kisiler.items():
        s = s.replace(takma, gercek)
    return s


def denetle(c: dict, metin: str, ibanlar: dict[str, str]) -> list[str]:
    """Modelin çıkarımındaki sorunlar (belgeyle ilgili bulgular değil). Boş liste = sorunsuz."""
    sorun = []
    for i, t in enumerate(c["talimatlar"], 1):
        on = f"Talimat {i}"
        for a in METIN_ALANLARI:
            v = t[a].strip()
            if v and not metinde_var(v, metin):
                sorun.append(f"{on}: {ETIKET[a]} '{v[:60]}' belgede geçmiyor (belgedeki yazımı aynen kullanın)")
        for a in ("gonderen_iban", "alici_iban"):
            v = t[a].strip()
            if v and v not in ibanlar:
                sorun.append(f"{on}: {ETIKET[a]} '{v[:40]}' belgedeki bir IBAN yer tutucusu değil ({', '.join(ibanlar) or 'belgede IBAN yok'}); "
                             "yalnız yer tutucuyu yazın ya da boş bırakın")
        if t["gonderen_iban"] and t["gonderen_iban"] == t["alici_iban"]:
            sorun.append(f"{on}: gönderen ve alıcı IBAN'ı aynı ({t['alici_iban']}); hangisinin alıcıya ait olduğunu yeniden okuyun")
        if not t["tutar"].strip():
            sorun.append(f"{on}: rakamla tutar boş; belgede rakamla tutar varsa yazın")
        elif para(t["tutar"]) is None:
            sorun.append(f"{on}: tutar '{t['tutar']}' sayı olarak okunamadı; yalnız rakamları belgedeki gibi yazın")
        for ad in t["imzalayanlar"]:
            if not metinde_var(ad, metin):
                sorun.append(f"{on}: imzalayan '{ad[:40]}' belgede geçmiyor")
    return sorun


def bulgular(t: dict, metin: str) -> list[tuple[str, str]]:
    """Belgeyle ilgili, kodun ürettiği bulgular (IBAN'lar geri açılmış talimat üzerinde). Model sonucu değiştirmez."""
    b = []
    for a, ad in (("gonderen_iban", "Gönderen IBAN"), ("alici_iban", "Alıcı IBAN")):
        if t[a] and (h := iban_kontrol(t[a])):
            b.append(("Hata", f"{ad}: {h}"))
    if not t["alici_iban"]:
        b.append(("Hata", "Alıcı IBAN'ı belgede bulunamadı"))
    if not t["gonderen_iban"]:
        b.append(("Dikkat", "Borçlanacak hesabın IBAN'ı belgede yok; müşteri hesabından eşleştirin"))
    rakam, yazi = para(t["tutar"]), yazidan_tutar(t["tutar_yazi"])
    if not t["tutar_yazi"]:
        b.append(("Dikkat", "Yazıyla tutar yok"))
    elif yazi is None:
        b.append(("Dikkat", f"Yazıyla tutar okunamadı: '{t['tutar_yazi']}'"))
    elif rakam is not None and yazi != rakam:
        b.append(("Hata", f"Rakamla tutar {tl(rakam)}, yazıyla tutar {tl(yazi)}: uyuşmuyor"))
    if not t["imzalayanlar"]:
        b.append(("Hata", "İmzalayan / gönderen kişi belgede bulunamadı"))
    if t["tur"] == "Belirtilmemiş":
        b.append(("Dikkat", "İşlem türü (EFT / FAST / havale / virman) belgede belirtilmemiş"))
    if t["para_birimi"] in ("Belirtilmemiş", "Diğer"):
        b.append(("Dikkat", f"Para birimi: {t['para_birimi']}"))
    if t["tarih"] and tarih_coz(t["tarih"]) is None:
        b.append(("Dikkat", f"Tarih okunamadı: '{t['tarih']}'"))
    k = katla(metin[:600])
    if re.search(r"\b(kimden|from|konu|subject)\b", k) and re.search(r"\b(kime|to|gonderildi|sent)\b", k):
        b.append(("Dikkat", "E-postayla gelmiş görünüyor: ıslak imza yok. Bankanızın e-posta / faks talimat prosedürüne göre "
                            "(sözleşme, geri arama teyidi) kontrol edin"))
    return b


def uret(metin: str, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    ek = ""
    if duzeltme:
        ek = ("\n<duzeltme>\nÖnceki çıkarım: " + json.dumps(onceki, ensure_ascii=False) + "\nKodun bulduğu sorunlar:\n"
              + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")
    return llm.json_iste(sistem, f"<talimat>\n{metin}\n</talimat>" + ek, SEMA)


def duzeltme_dongusu(uret_f, denetle_f, tur: int) -> tuple[dict, list[str], int]:
    taslak = uret_f(None, None)
    kalan = denetle_f(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret_f(kalan, taslak)
        kalan = denetle_f(taslak)
    return taslak, kalan, yapilan


def isle(yol: Path, kisiler: dict[str, str], tur: int) -> dict:
    """Bir belgeyi okur, maskeler, modele çıkarttırır, denetler; IBAN ve adları geri açılmış talimatları döndürür."""
    ham = belge.metin_oku(yol)
    if len(ham.strip()) < 20:
        raise belge.BelgeHatasi(f"{yol.name}: metin bulunamadı.")
    maskeli, ibanlar = maskele(ham, kisiler)
    c, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(maskeli, d, o), lambda x: denetle(x, maskeli, ibanlar), tur)
    talimatlar = []
    for i, t in enumerate(c["talimatlar"], 1):
        t = {a: (geri_ac(v, kisiler) if isinstance(v, str) else [geri_ac(x, kisiler) for x in v] if isinstance(v, list) else v) for a, v in t.items()}
        for a in ("gonderen_iban", "alici_iban"):
            t[a] = ibanlar.get(t[a].strip(), "")
        ref = t["referans"].strip()
        t["no"] = (f"{ref}/{i}" if len(c["talimatlar"]) > 1 else ref) if ref else f"{yol.stem}-{i}"
        t["dosya"] = yol.name
        t["bulgular"] = bulgular(t, ham)
        t["kalan"] = [s.split(": ", 1)[1] for s in kalan if s.startswith(f"Talimat {i}:")]
        onem = {x[0] for x in t["bulgular"]} | ({"Dikkat"} if t["kalan"] else set())
        t["durum"] = "Hatalı" if "Hata" in onem else "Kontrol" if "Dikkat" in onem else "Uygun"
        talimatlar.append(t)
    return {"dosya": yol.name, "talimatlar": talimatlar, "notlar": [geri_ac(n, kisiler) for n in c["notlar"]], "kalan": kalan,
            "tur": yapilan, "iban_sayisi": len(ibanlar), "karakter": len(maskeli)}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

KOD_BLOGU_SUTUNLARI = ["Talimat No", "Tarih", "Tür", "Müşteri No", "Gönderen Unvan", "Gönderen IBAN", "Alıcı Unvan", "Alıcı IBAN", "Tutar",
                       "Tutar (Yazı)", "Para Birimi", "Açıklama", "İmzalayanlar"]


def _baslik(ws, basliklar, genislik) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font, h.alignment = BASLIK_DOLGU, BASLIK_YAZI, UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def rapor_yaz(cikti: Path, sonuc: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Talimatlar"
    _baslik(ws, KOD_BLOGU_SUTUNLARI + ["Durum", "Bulgular", "Kaynak Dosya"], (16, 11, 10, 11, 28, 30, 28, 30, 14, 36, 8, 30, 24, 10, 60, 24))
    for d in sonuc:
        for t in d["talimatlar"]:
            pb = "" if t["para_birimi"] == "Belirtilmemiş" else t["para_birimi"]
            tur = "" if t["tur"] == "Belirtilmemiş" else t["tur"]
            notlar = [f"{o}: {m}" for o, m in t["bulgular"]] + [f"Kontrol (çıkarım): {k}" for k in t["kalan"]]
            ws.append([t["no"], t["tarih"], tur, "", t["gonderen_unvan"], t["gonderen_iban"], t["alici_unvan"], t["alici_iban"],
                       float(para(t["tutar"])) if para(t["tutar"]) is not None else t["tutar"], t["tutar_yazi"], pb, t["aciklama"],
                       ", ".join(t["imzalayanlar"]), t["durum"], "\n".join(notlar), d["dosya"]])
            n = ws.max_row
            ws.cell(n, 9).number_format = "#,##0.00"
            for j in (5, 6, 7, 8, 9, 10, 12, 13):
                ws.cell(n, j).fill = AI_DOLGU
            ws.cell(n, 14).fill = PatternFill("solid", fgColor=RENK[t["durum"]])
            for h in ws[n]:
                h.alignment = UST

    b = wb.create_sheet("Bulgular")
    _baslik(b, ["Talimat No", "Önem", "Bulgu", "Kaynak Dosya"], (16, 9, 90, 24))
    for d in sonuc:
        for t in d["talimatlar"]:
            for o, m in t["bulgular"]:
                b.append([t["no"], o, m, d["dosya"]])
                b.cell(b.max_row, 2).fill = PatternFill("solid", fgColor=RENK[o])
        for n in d["notlar"]:
            b.append(["—", "Bilgi", f"Model notu: {n}", d["dosya"]])
            b.cell(b.max_row, 2).fill = PatternFill("solid", fgColor=RENK["Bilgi"])
        if not d["talimatlar"]:
            b.append(["—", "Dikkat", "Bu belgede ödeme talimatı bulunamadı", d["dosya"]])
            b.cell(b.max_row, 2).fill = PatternFill("solid", fgColor=RENK["Dikkat"])

    k = wb.create_sheet("Kontrol")
    _baslik(k, ["Dosya", "Talimat", "IBAN (yer tutucu)", "Düzeltme Turu", "Son kontrol"], (28, 9, 12, 10, 90))
    for d in sonuc:
        k.append([d["dosya"], len(d["talimatlar"]), d["iban_sayisi"], d["tur"], "; ".join(d["kalan"]) or "Kod denetiminden sorunsuz geçti"])
        k.cell(k.max_row, 5).fill = KIRMIZI if d["kalan"] else YESIL
        k.cell(k.max_row, 5).alignment = UST
    k.append([])
    k.append([f"Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    k.append(["Önemli: Mor hücreler yapay zekâ ile çıkarılmıştır. İşlemi gerçekleştirmeden önce orijinal talimatla karşılaştırın; "
              "imza ve yetki kontrolü için 'Müşteri Talimat Kontrolü' kod bloğunu kullanın (Müşteri No sütununu doldurun)."])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def dosyalar(girdi: Path) -> list[Path]:
    if girdi.is_dir():
        l = sorted(p for p in girdi.iterdir() if p.is_file() and p.suffix.lower() in belge.DESTEKLENEN)
        if not l:
            raise belge.BelgeHatasi(f"{girdi}: klasörde .pdf, .docx veya .txt dosyası yok.")
        return l
    if not girdi.exists():
        raise belge.BelgeHatasi(f"{girdi}: dosya bulunamadı.")
    return [girdi]


def calistir(girdi: Path, cikti: Path, gizle: list[str] | None = None, tur: int = 2, evet: bool = False) -> dict:
    l = dosyalar(girdi)
    kisiler = takma_adlar(gizle or [])
    llm.onay_al(f"{len(l)} talimat belgesinin metni gönderilecek. IBAN'lar gönderilmez ([IBAN-n] yer tutucusu); telefon, e-posta, TCKN"
                f"{f' ve {len(kisiler)} ad' if kisiler else ''} maskeli. Unvanlar, adlar, tutarlar ve açıklamalar gönderilir.", evet)
    sonuc = [isle(y, kisiler, tur) for y in l]
    rapor_yaz(cikti, sonuc)
    js = cikti.with_suffix(".json")
    js.write_text(json.dumps(sonuc, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    return {"dosyalar": sonuc, "talimatlar": [t for d in sonuc for t in d["talimatlar"]], "json": js}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Yazılı ödeme talimatlarından gönderen, alıcı, IBAN, tutar ve açıklamayı çıkarır ve denetler.")
    p.add_argument("--girdi", type=Path, default=ORNEK, help="Talimat belgesi ya da klasörü (.pdf, .docx, .txt)")
    p.add_argument("--gizle", nargs="*", default=[], metavar="AD", help="Ayrıca maskelenecek kişi veya şirket adları")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "talimatlar.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.cikti, a.gizle, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    say = {d: sum(t["durum"] == d for t in s["talimatlar"]) for d in ("Hatalı", "Kontrol", "Uygun")}
    print(f"[OK] {len(s['dosyalar'])} belge · {len(s['talimatlar'])} talimat · {say['Hatalı']} hatalı, {say['Kontrol']} kontrol, {say['Uygun']} uygun")
    for t in s["talimatlar"]:
        for o, m in t["bulgular"]:
            if o == "Hata":
                print(f"[!] {t['no']} ({t['alici_unvan']}): {m}")
    print(f"[OK] Rapor: {a.cikti.resolve()} (+ {s['json'].name})")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
