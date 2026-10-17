"""
Basın Bülteni Taslağı — Gerisi Bende AI Agent
Pazarlama › Pazarlama Müdürü

1. Kod bilgi notunu okur: şirket, şehir, tarih, basın iletişimi ve "şirket hakkında" metni; numaralı bilgiler (B1, B2 …)
   ve konuşmacı adıyla alıntılar (A1, A2 …).
2. Model kurumsal dilde başlık, spot ve paragraflar yazar; her paragraf dayandığı bilgi / alıntı numaralarını gösterir.
3. Kod denetler:
   - Başlık, spot ve paragraflardaki sayılar bilgi notunda (paragraflarda dayandığı bilgilerde) yazıyor mu
   - Tırnak içindeki her söz verilen alıntılardan birinde aynen geçiyor mu, paragraf o alıntıya dayanıyor mu ve konuşmacının
     adı paragrafta yazıyor mu (alıntı uydurulmasın)
   - "İlk", "lider", "en büyük", "rekor" gibi iddialar bilgi notunda geçiyor mu
   - Başlık ve spot uzunluğu
   Sorun varsa modele geri gönderilir (en fazla --tur kez).
4. Kod yer-tarih satırını, "şirket hakkında" ve basın iletişimi bölümlerini aynen ekler.

Kullanım:
    python agent.py                                          # örnek (kurgusal) santral açılışı
    python agent.py --girdi bilgi_notu.txt --baslik-azami 90
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
SEMA = {
    "type": "object",
    "properties": {
        "baslik": S, "spot": S,
        "paragraflar": {"type": "array", "items": {"type": "object", "properties": {"metin": S, "dayanak": {"type": "array", "items": S}},
                                                   "required": ["metin", "dayanak"], "additionalProperties": False}},
    },
    "required": ["baslik", "spot", "paragraflar"],
    "additionalProperties": False,
}
ALANLAR = {"sirket": "Şirket", "sehir": "Şehir", "tarih": "Tarih", "iletisim": "Basın İletişimi", "hakkinda": "Şirket Hakkında"}
ALINTI = re.compile(r"^\s*(?:[-•*]\s*)?([^:\"“]+?),\s*([^:\"“]+?)\s*:\s*[\"“](.+)[\"”]\s*$")
TIRNAK = re.compile(r"[\"“”«»]([^\"“”«»]{8,})[\"“”«»]")
IDDIA = re.compile(r"\b(ilk|lider\w*|en buyuk\w*|en cok|en fazla|en iyi\w*|en genis\w*|tek\b|benzersiz|essiz|rekor\w*|oncu\w*|dunyanin|"
                   r"bir numara\w*|sektorun en)")


def notu_oku(yol: Path) -> dict:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: bilgi notu bulunamadı.")
    n = {a: "" for a in ALANLAR} | {"bilgiler": [], "alintilar": []}
    bolum = None
    for s in belge.metin_oku(yol).splitlines():
        if not s.strip():
            continue
        a, _, d = s.partition(":")
        anahtar = next((k for k, v in ALANLAR.items() if katla(v) == katla(a)), None)
        if anahtar and d.strip():
            n[anahtar], bolum = d.strip(), None
        elif katla(s) in {"bilgiler", "alintilar"}:
            bolum = katla(s)
        elif bolum == "alintilar" and ALINTI.match(s):
            m = ALINTI.match(s)
            n["alintilar"].append({"id": f"A{len(n['alintilar']) + 1}", "ad": m.group(1).strip(), "unvan": m.group(2).strip(), "soz": m.group(3).strip()})
        elif bolum == "bilgiler":
            n["bilgiler"].append(s.strip().lstrip("-•* ").strip())
    if not n["sirket"] or not n["bilgiler"]:
        raise llm.LLMHatasi(f"{yol.name}: 'Şirket: …' satırı ve 'Bilgiler' bölümü gerekli.")
    return n


def denetle(t: dict, n: dict, baslik_azami: int, spot_azami: int) -> list[str]:
    sorun = []
    b = {f"B{i}": s for i, s in enumerate(n["bilgiler"], 1)}
    a = {x["id"]: x for x in n["alintilar"]}
    tum_bilgi = " ".join(n["bilgiler"]) + " " + n["tarih"]
    tum_kaynak = katla(tum_bilgi + " " + " ".join(x["soz"] for x in n["alintilar"]) + " " + n["hakkinda"])
    for ad, metin, sinir in (("Başlık", t["baslik"], baslik_azami), ("Spot", t["spot"], spot_azami)):
        if len(metin) > sinir:
            sorun.append(f"{ad}: {len(metin)} karakter, sınır {sinir}")
        fazla = sorted(sayilar(metin) - sayilar(tum_bilgi))
        if fazla:
            sorun.append(f"{ad}: bilgi notunda olmayan sayı: {', '.join(fazla)}")
    parcalar = [("Başlık", t["baslik"], None), ("Spot", t["spot"], None)] + [(f"Paragraf {i}", p["metin"], p["dayanak"]) for i, p in enumerate(t["paragraflar"], 1)]
    for ad, metin, dayanak in parcalar:
        for m in IDDIA.finditer(katla(metin)):
            if m.group(0) not in tum_kaynak:
                sorun.append(f"{ad}: '{m.group(0)}' iddiası bilgi notunda yok")
        if dayanak is None:
            if TIRNAK.search(metin):
                sorun.append(f"{ad}: başlık ve spotta alıntı kullanmayın")
            continue
        gec = [d for d in dayanak if d not in b and d not in a]
        if not dayanak or gec:
            sorun.append(f"{ad}: dayanak yok ya da geçersiz {gec}")
        kaynak = " ".join(b.get(d, "") for d in dayanak) + " " + n["tarih"] + " " + " ".join(a[d]["soz"] for d in dayanak if d in a)
        fazla = sorted(sayilar(TIRNAK.sub(" ", metin)) - sayilar(kaynak))
        if fazla:
            sorun.append(f"{ad}: dayandığı bilgilerde olmayan sayı: {', '.join(fazla)}")
        for q in TIRNAK.findall(metin):
            kim = next((x for x in n["alintilar"] if metinde_var(q, x["soz"])), None)
            if not kim:
                sorun.append(f"{ad}: tırnak içindeki söz verilen alıntılarda yok: '{q[:60]}' (alıntı uydurulmamalı)")
            elif kim["id"] not in dayanak:
                sorun.append(f"{ad}: {kim['id']} alıntısı kullanılmış ama dayanakta yok")
            elif katla(kim["ad"].split()[-1]) not in katla(metin):
                sorun.append(f"{ad}: {kim['id']} alıntısının sahibi ({kim['ad']}) paragrafta belirtilmemiş")
    if not t["paragraflar"]:
        sorun.append("Genel: paragraf yok")
    return sorun


def uret(n: dict, baslik_azami: int, spot_azami: int, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    b = "\n".join(f"B{i} | {s}" for i, s in enumerate(n["bilgiler"], 1))
    a = "\n".join(f"{x['id']} | {x['ad']}, {x['unvan']}: \"{x['soz']}\"" for x in n["alintilar"]) or "(alıntı yok)"
    mesaj = (f"<sirket>{n['sirket']}</sirket>\n<bilgiler>\n{b}\n</bilgiler>\n<alintilar>\n{a}\n</alintilar>\n"
             f"<sinirlar baslik=\"{baslik_azami}\" spot=\"{spot_azami}\"/>" + duzeltme_blogu(duzeltme, onceki))
    return llm.json_iste(sistem, mesaj, SEMA)


def bulten(t: dict, n: dict) -> str:
    yer = ", ".join(x for x in (n["sehir"], n["tarih"]) if x)
    p = [x["metin"] for x in t["paragraflar"]]
    s = ["BASIN BÜLTENİ", "", t["baslik"].upper(), "", t["spot"], "", (f"{yer} – " if yer else "") + p[0] if p else ""]
    s += [x for y in p[1:] for x in ("", y)]
    if n["hakkinda"]:
        s += ["", f"{n['sirket']} Hakkında", n["hakkinda"]]
    if n["iletisim"]:
        s += ["", "Basın İletişimi", n["iletisim"]]
    return "\n".join(s)


def rapor_yaz(cikti: Path, t, n, kalan, tur, metin) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Bülten"
    baslik(ws, ["Bölüm", "Metin", "Dayanak", "Dayanak Metni"], (12, 90, 12, 70))
    ws.append(["Başlık", t["baslik"], "", ""])
    ws.append(["Spot", t["spot"], "", ""])
    b = {f"B{i}": s for i, s in enumerate(n["bilgiler"], 1)} | {x["id"]: f"{x['ad']}: {x['soz']}" for x in n["alintilar"]}
    for i, p in enumerate(t["paragraflar"], 1):
        ws.append([f"Paragraf {i}", p["metin"], ", ".join(p["dayanak"]), "\n".join(f"{d}. {b.get(d, '?')}" for d in p["dayanak"])])
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 2).fill = AI_DOLGU
    son_not(ws, "Önemli: Bülten taslaktır. Rakamları, alıntıları ve iddiaları ilgili birimlere ve konuşmacılara onaylatmadan yayımlamayın.", kalan, tur)
    hizala(ws)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    cikti.with_suffix(".txt").write_text(metin, encoding="utf-8")


def calistir(girdi: Path, cikti: Path, baslik_azami: int = 100, spot_azami: int = 300, tur: int = 2, evet: bool = False) -> dict:
    n = notu_oku(girdi)
    print(f"[OK] {n['sirket']} · {len(n['bilgiler'])} bilgi · {len(n['alintilar'])} alıntı")
    llm.onay_al(f"Bilgi notu ({len(n['bilgiler'])} bilgi, {len(n['alintilar'])} alıntı) gönderilecek; basın iletişimi bilgileri gönderilmez.", evet)
    t, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(n, baslik_azami, spot_azami, d, o), lambda x: denetle(x, n, baslik_azami, spot_azami), tur)
    metin = bulten(t, n)
    rapor_yaz(cikti, t, n, kalan, yapilan, metin)
    return {"not": n, "taslak": t, "kalan": kalan, "tur": yapilan, "metin": metin}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    p = argparse.ArgumentParser(description="Duyuru bilgisinden kurumsal dilde basın bülteni taslağı hazırlar.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "bilgi_notu.txt",
                   help="Bilgi notu (.txt/.docx): Şirket, Şehir, Tarih, Basın İletişimi, Şirket Hakkında satırları; 'Bilgiler' ve 'Alıntılar' bölümleri")
    p.add_argument("--baslik-azami", type=int, default=100, help="Başlık için en fazla karakter (varsayılan 100)")
    p.add_argument("--spot-azami", type=int, default=300, help="Spot için en fazla karakter (varsayılan 300)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "basin_bulteni.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.cikti, a.baslik_azami, a.spot_azami, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {s['taslak']['baslik']}")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Bülten: {a.cikti.with_suffix('.txt').resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
