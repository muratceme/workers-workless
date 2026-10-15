"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

NOT = KLASOR / "ornek_veri" / "bilgi_notu.txt"
MESAJLAR = []


def p(metin, *d):
    return {"metin": metin, "dayanak": list(d)}


def taslak(dogru=True) -> dict:
    t = {"baslik": "Örnek Enerji, Soma'da 45 MW gücündeki güneş enerjisi santralini devreye aldı",
         "spot": "Örnek Enerji A.Ş., Manisa'nın Soma ilçesinde 38 milyon ABD doları yatırımla kurduğu 45 MW kurulu güce sahip güneş enerjisi "
                 "santralini devreye aldı.",
         "paragraflar": [p("Örnek Enerji, Manisa'nın Soma ilçesinde kurulu gücü 45 MW olan güneş enerjisi santralini devreye aldı. Yatırım tutarı "
                           "38 milyon ABD doları olan santralin yılda yaklaşık 80 milyon kWh elektrik üretmesi bekleniyor.", "B1", "B2", "B3"),
                         p("Üretilecek elektrik yaklaşık 30 bin hanenin yıllık tüketimine denk geliyor. Bu yatırımla şirketin toplam kurulu gücü "
                           "310 MW'a ulaştı.", "B5", "B6"),
                         p("Santralin yapımında 250 kişi çalıştı; işletme döneminde 20 kişi istihdam edilecek. Genel Müdür Ayşe Kara, “Bölgede "
                           "istihdama katkı sunmaktan mutluluk duyuyoruz.” dedi.", "B4", "A1")]}
    if not dogru:
        t["baslik"] = "Örnek Enerji'den Türkiye'nin en büyük güneş santrali"
        t["paragraflar"][0]["metin"] = t["paragraflar"][0]["metin"].replace("45 MW", "50 MW")
        t["paragraflar"][2]["metin"] += " Kara, “Bu yatırım sektörde bir dönüm noktasıdır.” diye ekledi."
    return t


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return taslak("<duzeltme>" in kullanici)


class KodTesti(unittest.TestCase):
    def test_okuma(self):
        n = agent.notu_oku(NOT)
        self.assertEqual((n["sirket"], n["sehir"], n["tarih"], len(n["bilgiler"])), ("Örnek Enerji A.Ş.", "İzmir", "15.10.2026", 6))
        self.assertEqual((n["alintilar"][0]["ad"], n["alintilar"][0]["unvan"]), ("Ayşe Kara", "Genel Müdür"))

    def test_denetim(self):
        n = agent.notu_oku(NOT)
        self.assertEqual(agent.denetle(taslak(), n, 100, 300), [])
        s = agent.denetle(taslak(False), n, 100, 300)
        self.assertEqual(len(s), 3, s)
        self.assertTrue(any("en buyuk" in x for x in s) and any("50" in x for x in s) and any("dönüm noktası" in x for x in s))
        self.assertTrue(any("Başlık" in x and "sınır 50" in x for x in agent.denetle(taslak(), n, 50, 300)))
        t = taslak()
        t["paragraflar"][2]["dayanak"] = ["B4"]
        self.assertTrue(any("A1 alıntısı" in x for x in agent.denetle(t, n, 100, 300)))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "b.xlsx"
            s = agent.calistir(NOT, cikti, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            self.assertNotIn("basin@ornekenerji", MESAJLAR[0])
            m = cikti.with_suffix(".txt").read_text(encoding="utf-8")
            self.assertIn("İzmir, 15.10.2026 – Örnek Enerji, Manisa'nın", m)
            self.assertIn("Örnek Enerji A.Ş. Hakkında", m)
            self.assertTrue(m.rstrip().endswith("0232 000 00 00"))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Bülten"])

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(NOT, Path(tmp) / "b.xlsx", tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 3)
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--evet", "--girdi", str(NOT.with_name("yok.txt"))]), 1)


if __name__ == "__main__":
    unittest.main()
