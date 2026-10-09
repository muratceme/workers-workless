"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ANKET, SORULAR = agent.ORNEK / "anket.csv", agent.ORNEK / "sorular.csv"
MESAJLAR = []
TEMA = {  # id → (temalar, duygu, risk, alıntı)
    "A001": (["Ücret ve yan haklar"], "olumsuz", "yok", "Ücretler enflasyon karşısında eridi"),
    "A002": (["Yönetici ve liderlik", "Ekip ve iş arkadaşları"], "olumlu", "yok", "Ustabaşımız çok destekleyici"),
    "A003": (["Vardiya ve çalışma saatleri"], "olumsuz", "yok", "Vardiya değişimleri son dakika bildiriliyor"),
    "A004": (["Çalışma koşulları"], "olumsuz", "yok", "Yemekhane yemekleri son aylarda kötüleşti"),
    "A005": (["İş sağlığı ve güvenliği"], "olumsuz", "iş sağlığı ve güvenliği", "Eğitim verilmeden yeni makinede çalışmaya başlatıldık"),
    "A006": (["Takdir ve adalet", "Ücret ve yan haklar"], "olumsuz", "yok", "Prim sistemi şeffaf değil"),
    "A019": (["Yönetici ve liderlik"], "olumsuz", "mobbing / taciz / kötü muamele", "herkesin önünde azarlanıyoruz"),
    "A020": (["İş yükü ve denge", "Vardiya ve çalışma saatleri"], "olumsuz", "yok", "hafta sonları da çağrılıyoruz"),
    "A021": (["İş yükü ve denge"], "olumsuz", "yok", "iş yükü çok arttı ve herkes yoruldu"),          # yorumda yok → çıkarılmalı
    "A022": (["Yönetici ve liderlik", "İletişim"], "olumsuz", "mobbing / taciz / kötü muamele", "söyleyen hedef haline geliyor"),
    "A023": (["İş sağlığı ve güvenliği"], "olumsuz", "iş sağlığı ve güvenliği", "frenleri zayıf olan araçla çalışıyoruz"),
    "A029": (["Ücret ve yan haklar"], "olumsuz", "yok", "sabit maaş piyasanın altında"),
    "A030": (["Gelişim ve kariyer", "Yönetici ve liderlik"], "olumlu", "yok", "Yöneticim gelişimime gerçekten önem veriyor"),
    "A031": (["Ücret ve yan haklar"], "olumsuz", "yok", "Yol ve yemek ücreti yetersiz"),
    "A038": (["İş yükü ve denge"], "olumsuz", "yok", "sürekli geç saatlere kalıyoruz"),
    "A039": (["İletişim"], "olumlu", "yok", "Ekip içi iletişim iyi"),
    "A044": (["Çalışma koşulları"], "olumlu", "yok", "Uzaktan çalışma imkânı çok değerli"),
}


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append((sistem, kullanici))
    if "<yorumlar>" in kullanici:
        idler = re.findall(r'<yorum id="([^"]+)"', kullanici)
        return {"yorumlar": [{"id": i, "temalar": TEMA[i][0] + ["Uydurma tema"], "duygu": TEMA[i][1], "oneri": "", "risk_turu": TEMA[i][2],
                              "alinti": TEMA[i][3]} for i in idler]
                + [{"id": "A999", "temalar": ["Diğer"], "duygu": "notr", "oneri": "", "risk_turu": "yok", "alinti": ""}]}
    return {"basliklar": [
        {"baslik": "Ücretin yetersiz bulunması", "aciklama": "...", "temalar": ["Ücret ve yan haklar"], "yorum_idleri": ["A001", "A029", "A031", "A999"],
         "olasi_aksiyon": "Ücret araştırması ile karşılaştırma"},
        {"baslik": "Yönetici davranışı", "aciklama": "...", "temalar": ["Yönetici ve liderlik"], "yorum_idleri": ["A019", "A022"], "olasi_aksiyon": "..."},
        {"baslik": "Dayanaksız başlık", "aciklama": "...", "temalar": ["Diğer"], "yorum_idleri": ["X1"], "olasi_aksiyon": "..."}]}


class PuanTesti(unittest.TestCase):
    def test_puan_analizi_yorumsuz(self):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(ANKET, SORULAR, Path(tmp) / "b.xlsx", yorum=False, katilim={"Üretim": 24, "Depo": 12, "Bilgi İşlem": 4, "Lojistik": 5})
        p = s["puan"]
        self.assertEqual(list(p["birim"]), ["Depo", "Muhasebe", "Satış", "Üretim"])        # Bilgi İşlem (3 yanıt) gizli
        self.assertEqual(round(p["sirket_enps"]["enps"]), -28)
        self.assertAlmostEqual(p["sirket"]["Ücret ve Yan Haklar"]["olumlu"], 15 / 92)
        self.assertEqual(round(p["birim"]["Depo"]["enps"]["enps"]), -90)
        self.assertEqual(s["katilim"], {"Üretim": 24, "Depo": 12, agent.DIGER: 4})
        self.assertEqual({y.id for y in s["yorumlu"] if y.hassas}, {"A005", "A019", "A022", "A023"})   # anahtar kelime taraması
        turler = {(u["tur"], u["kim"]) for u in s["uyarilar"]}
        self.assertTrue({("Birimde düşük skor", "Depo"), ("Şirket geneli zayıf alan", ""), ("Katılım", "Lojistik")} <= turler)
        self.assertEqual(MESAJLAR, [])

    def test_ters_madde_ve_gecersiz(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t / "s.csv").write_text("Soru Kodu;Soru;Boyut;Ters Madde\nQ1;x;Yük;Evet\nQ2;y;eNPS;Hayır\n", encoding="utf-8")
            (t / "a.csv").write_text("Birim;Q1;Q2;Yorum\n" + "X;5;10;\n" * 4 + "X;1;0;Ali Bey bize bağırıyor, tel 0532 111 22 33\nX;9;11;\n",
                                     encoding="utf-8")
            s = agent.calistir(t / "a.csv", t / "s.csv", t / "o.xlsx", min_yanit=5, yorum=False)
            p = s["puan"]
            self.assertAlmostEqual(p["sirket"]["Yük"]["ort"], 1.8)                       # 4 × (6−5) + (6−1) = 9 / 5
            self.assertEqual(p["sirket_enps"]["enps"], 60)                               # 4 destekçi, 1 kötüleyen / 5
            self.assertTrue(any(u["tur"] == "Geçersiz puan" for u in s["uyarilar"]))
            self.assertEqual(s["yorumlu"][0].maskeli, "[KİŞİ] Bey bize bağırıyor, tel [TELEFON]")


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_uctan_uca(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "b.xlsx"
            s = agent.calistir(ANKET, SORULAR, cikti, evet=True)
            self.assertEqual(len(MESAJLAR), 2)
            self.assertNotIn("Depo", MESAJLAR[0][1])                                     # birim bilgisi gönderilmez
            self.assertIn("## Yorum analizi", MESAJLAR[0][0])
            self.assertNotIn("## Sorun başlıkları", MESAJLAR[0][0])
            self.assertIn("## Sorun başlıkları", MESAJLAR[1][0])
            y = {x.id: x for x in s["yorumlu"]}
            self.assertEqual(y["A021"].analiz["alinti"], "")
            self.assertIn("Alıntı yorumda bulunamadı; çıkarıldı", y["A021"].kontrol)
            self.assertEqual(y["A001"].analiz["temalar"], ["Ücret ve yan haklar"])
            self.assertEqual(y["A019"].hassas, "mobbing / taciz / kötü muamele")
            self.assertEqual(s["temalar"]["Ücret ve yan haklar"]["toplam"], 4)
            self.assertEqual(dict(s["temalar"]["Çalışma koşulları"]["birim"]), {"Üretim": 1, "Gizli": 1})
            self.assertEqual([b["baslik"] for b in s["basliklar"]], ["Ücretin yetersiz bulunması", "Yönetici davranışı"])
            self.assertEqual(s["basliklar"][0]["yorum_idleri"], ["A001", "A029", "A031"])
            self.assertTrue(any(u["tur"] == "Başlık çıkarıldı" for u in s["uyarilar"]))
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Özet", "Hassas Yorumlar", "Birim × Boyut", "Sorular", "Sorun Başlıkları", "Temalar", "Yorumlar", "Uyarılar"])
            self.assertEqual(wb["Sorun Başlıkları"]["D2"].value, 3)
            self.assertEqual(wb["Hassas Yorumlar"].max_row, 7)
            self.assertEqual(wb["Yorumlar"]["B18"].value, "Gizli")

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(agent.main(["--yorum-yok", "--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(agent.main(["--calisan-sayisi", "Depo"]), 2)
            self.assertEqual(agent.main(["--anket", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
