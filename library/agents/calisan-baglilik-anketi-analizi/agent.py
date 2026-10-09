"""
Çalışan Bağlılık Anketi Analizi — Workers / Workless AI Agent
İnsan Kaynakları › İnsan Kaynakları Müdürü

1. Kod anket puanlarını analiz eder:
   - Likert (1–5) sorular boyutlara toplanır; ters maddeler çevrilir (6 − puan).
   - Boyut ve birim bazında ortalama, olumlu (4–5) ve olumsuz (1–2) oranı; 0–10 tavsiye sorusundan eNPS.
   - Anonimlik: yanıt sayısı --min-yanit altındaki birimler "Diğer (küçük birimler)" altında birleştirilir.
   - İsteğe bağlı katılım oranı (birim çalışan sayıları verilirse).
2. Model açık uçlu yorumları (kişisel veriler maskelenmiş, birim bilgisi gönderilmeden) temalara ayırır,
   duygu ve öneriyi çıkarır, hassas durumları (mobbing/taciz, ayrımcılık, iş güvenliği, etik) işaretler ve her yorum
   için yorumdan birebir kısa alıntı verir. Kod alıntının yorumda geçtiğini doğrular; ayrıca anahtar kelime ile
   hassas yorum taraması yapar (modelin kaçırdığı da insan incelemesine düşer).
3. Model, kodun saydığı tema özetinden temel sorun başlıklarını yazar; her başlık yorum kimliklerine dayanmalıdır.
   Kimliği doğrulanamayan başlık rapora alınmaz. Sayılar koddan gelir.

Kullanım:
    python agent.py                                           # örnek: 46 yanıt, 5 birim, 17 yorum
    python agent.py --anket anket.xlsx --sorular sorular.xlsx --min-yanit 5
    python agent.py --anket anket.xlsx --sorular sorular.xlsx --calisan-sayisi Üretim=25 Depo=12 --yorum-yok
"""
from __future__ import annotations

import argparse
import csv
import re
import statistics
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import llm

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
PAKET = 40
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
DIGER = "Diğer (küçük birimler)"
TEMALAR = ["Yönetici ve liderlik", "Ücret ve yan haklar", "Gelişim ve kariyer", "İş yükü ve denge", "Vardiya ve çalışma saatleri", "İletişim",
           "Çalışma koşulları", "İş sağlığı ve güvenliği", "Takdir ve adalet", "Ekip ve iş arkadaşları", "Diğer"]
RISK_TURLERI = ["yok", "mobbing / taciz / kötü muamele", "ayrımcılık", "iş sağlığı ve güvenliği", "etik / usulsüzlük", "şiddet veya tehdit",
                "kendine zarar"]
# Model kaçırsa da insan incelemesine düşmesi için anahtar kelimeler (katlanmış, kök hâlinde)
RISK_KELIMELERI = ["bagir", "azarla", "hakaret", "mobbing", "taciz", "tehdit", "asagila", "ayrimcilik", "irkci", "kaza", "yaralan", "fren",
                   "is guvenligi", "rusvet", "yolsuzluk", "intihar", "siddet", "hedef haline"]
YORUM_BASLIKLARI = ("gorus", "yorum", "oneri", "aciklama", "eklemek")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def maskele(s: str) -> str:
    """llm.maskele + unvanlı adlar (Ahmet Bey, Ayşe Hanım, Mehmet Usta)."""
    s = llm.maskele(s)
    return re.sub(r"\b[A-ZÇĞİÖŞÜ][a-zçğıöşü]+(?:\s+[A-ZÇĞİÖŞÜ][a-zçğıöşü]+)?\s+(Bey|Hanım|Hanim|Usta|Şef|Abi|Abla)\b", r"[KİŞİ] \1", s)


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Soru:
    kod: str
    metin: str
    boyut: str
    ters: bool

    @property
    def enps(self) -> bool:
        return katla(self.boyut) in ("enps", "tavsiye", "nps")


@dataclass
class Yanit:
    id: str
    birim: str
    grup: str
    puanlar: dict
    yorum: str
    rapor_birimi: str = ""
    maskeli: str = ""
    hassas: str = ""
    analiz: dict = field(default_factory=dict)
    kontrol: list = field(default_factory=list)


def oku_sorular(yol: Path) -> list[Soru]:
    s = tablo_oku(yol)
    b = [katla(x) for x in s[0]]
    al = lambda *a: next((b.index(x) for x in a if x in b), None)  # noqa: E731
    k, m, bo, t = al("soru kodu", "kod"), al("soru", "soru metni"), al("boyut", "kategori"), al("ters madde", "ters")
    if k is None or bo is None:
        raise ValueError(f"{yol.name}: 'Soru Kodu' ve 'Boyut' sütunları gerekli")
    return [Soru(metin(r[k]), metin(r[m]) if m is not None else "", metin(r[bo]), t is not None and katla(r[t]).startswith(("evet", "e", "1", "x")))
            for r in s[1:] if len(r) > k and metin(r[k])]


def oku_anket(yol: Path, sorular: list[Soru]) -> tuple[list[Yanit], list[dict]]:
    s = tablo_oku(yol)
    b = [katla(x) for x in s[0]]
    kodlar = {katla(q.kod): q for q in sorular}
    id_i = next((i for i, x in enumerate(b) if x in ("yanit no", "yanit id", "id", "no")), None)
    birim_i = next((i for i, x in enumerate(b) if x in ("birim", "departman", "bolum")), None)
    if birim_i is None:
        raise ValueError(f"{yol.name}: 'Birim' sütunu gerekli")
    grup_i = next((i for i, x in enumerate(b) if x in ("kidem", "kidem grubu", "grup", "yas grubu")), None)
    soru_i = {i: kodlar[x] for i, x in enumerate(b) if x in kodlar}
    yorum_i = [i for i, x in enumerate(b) if i not in soru_i and any(x.startswith(y) or y in x for y in YORUM_BASLIKLARI)]
    eksik = [q.kod for q in sorular if q not in soru_i.values()]
    uyarilar = [{"onem": "Orta", "tur": "Soru sütunu yok", "kim": ", ".join(eksik), "aciklama": "Soru listesindeki kodlar ankette bulunamadı"}] if eksik else []
    gecersiz = 0
    yanitlar = []
    for n, r in enumerate(s[1:], 1):
        r = r + [None] * (len(b) - len(r))
        puanlar = {}
        for i, q in soru_i.items():
            v = r[i]
            if v in (None, ""):
                continue
            try:
                v = float(str(v).replace(",", "."))
            except ValueError:
                gecersiz += 1
                continue
            if (q.enps and 0 <= v <= 10) or (not q.enps and 1 <= v <= 5):
                puanlar[q.kod] = (6 - v) if q.ters and not q.enps else v
            else:
                gecersiz += 1
        yorum = "\n".join(metin(r[i]) for i in yorum_i if metin(r[i]))
        yanitlar.append(Yanit(metin(r[id_i]) if id_i is not None and metin(r[id_i]) else f"Y{n:03d}", metin(r[birim_i]) or "—",
                              metin(r[grup_i]) if grup_i is not None else "", puanlar, yorum))
    if gecersiz:
        uyarilar.append({"onem": "Orta", "tur": "Geçersiz puan", "kim": "", "aciklama": f"{gecersiz} hücre ölçek dışı veya sayı değil; hesaba alınmadı"})
    if not yorum_i:
        uyarilar.append({"onem": "Bilgi", "tur": "Yorum sütunu yok", "kim": "", "aciklama": "Başlığında 'görüş', 'yorum' veya 'öneri' geçen sütun bulunamadı"})
    return yanitlar, uyarilar


# ----------------------------------------------------------------------------
# Puan analizi (kod)
# ----------------------------------------------------------------------------

def anonimlestir(yanitlar: list[Yanit], min_yanit: int) -> list[dict]:
    say = Counter(y.birim for y in yanitlar)
    kucuk = {b for b, n in say.items() if n < min_yanit}
    for y in yanitlar:
        y.rapor_birimi = DIGER if y.birim in kucuk else y.birim
    uy = []
    if kucuk:
        n = sum(say[b] for b in kucuk)
        uy.append({"onem": "Bilgi", "tur": "Anonimlik", "kim": "", "aciklama": f"{len(kucuk)} birim ({n} yanıt) {min_yanit} yanıtın altında; "
                   f"'{DIGER}' altında birleştirildi" + (" ve bu grup da küçük olduğu için birim tablolarında gösterilmez" if n < min_yanit else "")})
    return uy


def ozetle(degerler: list[float]) -> dict:
    if not degerler:
        return {"n": 0, "ort": None, "olumlu": None, "olumsuz": None}
    return {"n": len(degerler), "ort": statistics.fmean(degerler), "olumlu": sum(1 for v in degerler if v >= 4) / len(degerler),
            "olumsuz": sum(1 for v in degerler if v <= 2) / len(degerler)}


def enps(degerler: list[float]) -> dict:
    if not degerler:
        return {"n": 0, "enps": None, "destekci": 0, "notr": 0, "kotuleyen": 0}
    d = sum(1 for v in degerler if v >= 9)
    k = sum(1 for v in degerler if v <= 6)
    return {"n": len(degerler), "enps": (d - k) / len(degerler) * 100, "destekci": d, "notr": len(degerler) - d - k, "kotuleyen": k}


def puan_analizi(yanitlar: list[Yanit], sorular: list[Soru], min_yanit: int, fark_esik: float = 0.15) -> dict:
    boyutlar = list(dict.fromkeys(q.boyut for q in sorular if not q.enps))
    enps_kod = [q.kod for q in sorular if q.enps]

    def boyut_degerleri(lst, boyut):
        return [y.puanlar[q.kod] for y in lst for q in sorular if q.boyut == boyut and q.kod in y.puanlar]

    sirket = {b: ozetle(boyut_degerleri(yanitlar, b)) for b in boyutlar}
    sirket_enps = enps([y.puanlar[k] for y in yanitlar for k in enps_kod if k in y.puanlar])
    birimler = defaultdict(list)
    for y in yanitlar:
        birimler[y.rapor_birimi].append(y)
    birim = {}
    for ad, lst in sorted(birimler.items(), key=lambda i: (i[0] == DIGER, i[0])):
        if len(lst) < min_yanit:
            continue
        birim[ad] = {"n": len(lst), "boyut": {b: ozetle(boyut_degerleri(lst, b)) for b in boyutlar},
                     "enps": enps([y.puanlar[k] for y in lst for k in enps_kod if k in y.puanlar])}
    soru = [{"soru": q, **(enps([y.puanlar[q.kod] for y in yanitlar if q.kod in y.puanlar]) if q.enps
                           else ozetle([y.puanlar[q.kod] for y in yanitlar if q.kod in y.puanlar]))} for q in sorular]
    uyarilar = []
    for ad, x in birim.items():
        for b, o in x["boyut"].items():
            s = sirket[b]
            if o["olumlu"] is not None and s["olumlu"] is not None and s["olumlu"] - o["olumlu"] >= fark_esik:
                uyarilar.append({"onem": "Orta", "tur": "Birimde düşük skor", "kim": ad, "aciklama": f"{b}: olumlu %{o['olumlu'] * 100:.0f}, şirket "
                                 f"%{s['olumlu'] * 100:.0f} ({(s['olumlu'] - o['olumlu']) * 100:.0f} puan altında; {x['n']} kişi)"})
    for b, s in sirket.items():
        if s["olumlu"] is not None and s["olumlu"] < 0.4:
            uyarilar.append({"onem": "Orta", "tur": "Şirket geneli zayıf alan", "kim": "", "aciklama": f"{b}: olumlu yanıt %{s['olumlu'] * 100:.0f}, "
                             f"olumsuz %{s['olumsuz'] * 100:.0f}"})
    return {"boyutlar": boyutlar, "sirket": sirket, "sirket_enps": sirket_enps, "birim": birim, "soru": soru, "uyarilar": uyarilar}


# ----------------------------------------------------------------------------
# Yorum analizi (model)
# ----------------------------------------------------------------------------

SEMA_YORUM = {
    "type": "object",
    "properties": {"yorumlar": {"type": "array", "items": {
        "type": "object",
        "properties": {
            "id": {"type": "string"},
            "temalar": {"type": "array", "items": {"type": "string", "enum": TEMALAR}},
            "duygu": {"type": "string", "enum": ["olumlu", "olumsuz", "karma", "notr"]},
            "oneri": {"type": "string"},
            "risk_turu": {"type": "string", "enum": RISK_TURLERI},
            "alinti": {"type": "string"},
        },
        "required": ["id", "temalar", "duygu", "oneri", "risk_turu", "alinti"], "additionalProperties": False}}},
    "required": ["yorumlar"], "additionalProperties": False,
}
SEMA_BASLIK = {
    "type": "object",
    "properties": {"basliklar": {"type": "array", "items": {
        "type": "object",
        "properties": {"baslik": {"type": "string"}, "aciklama": {"type": "string"}, "temalar": {"type": "array", "items": {"type": "string", "enum": TEMALAR}},
                       "yorum_idleri": {"type": "array", "items": {"type": "string"}}, "olasi_aksiyon": {"type": "string"}},
        "required": ["baslik", "aciklama", "temalar", "yorum_idleri", "olasi_aksiyon"], "additionalProperties": False}}},
    "required": ["basliklar"], "additionalProperties": False,
}


def _normal(s: str) -> str:
    return " ".join(str(s).replace("İ", "i").replace("I", "ı").lower().split()).strip(" .,;:!?\"'“”‘’")


def istem(bolum: str) -> str:
    """prompt.md: ortak giriş + '## Yorum analizi' veya '## Sorun başlıkları' bölümü."""
    t = (BURASI / "prompt.md").read_text(encoding="utf-8")
    giris, _, kalan = t.partition("## Yorum analizi")
    yorum, _, baslik = kalan.partition("## Sorun başlıkları")
    return giris + ("## Yorum analizi" + yorum if bolum == "yorum" else "## Sorun başlıkları" + baslik)


def alinti_dogru(alinti: str, kaynak: str) -> bool:
    a = _normal(alinti)
    return bool(a) and len(a) >= 8 and a in _normal(kaynak)


def yorumlari_analiz_et(yorumlu: list[Yanit]) -> None:
    sistem = istem("yorum")
    for i in range(0, len(yorumlu), PAKET):
        parca = yorumlu[i:i + PAKET]
        mesaj = "\n".join(["<tema_listesi>" + "; ".join(TEMALAR) + "</tema_listesi>", "<yorumlar>",
                           *[f'<yorum id="{y.id}">\n{y.maskeli}\n</yorum>' for y in parca], "</yorumlar>"])
        yanit = llm.json_iste(sistem, mesaj, SEMA_YORUM)
        harita = {y.id: y for y in parca}
        for s in yanit.get("yorumlar", []):
            y = harita.get(s.get("id"))
            if y is None or y.analiz:
                continue
            s["temalar"] = [t for t in dict.fromkeys(s.get("temalar", [])) if t in TEMALAR] or ["Diğer"]
            if s.get("alinti") and not alinti_dogru(s["alinti"], y.maskeli):
                y.kontrol.append("Alıntı yorumda bulunamadı; çıkarıldı")
                s["alinti"] = ""
            y.analiz = s
    for y in yorumlu:
        if not y.analiz:
            y.kontrol.append("Model bu yorum için sonuç döndürmedi")


def sorun_basliklari(yorumlu: list[Yanit], tema_tablosu: dict) -> tuple[list[dict], list[str]]:
    analizli = [y for y in yorumlu if y.analiz]
    if not analizli:
        return [], []
    sistem = istem("baslik")
    satirlar = ["<tema_ozeti>"]
    for t, x in tema_tablosu.items():
        satirlar.append(f'<tema ad="{t}" yorum="{x["toplam"]}" olumsuz="{x["olumsuz"]}">')
        for y in [y for y in analizli if t in y.analiz["temalar"]][:12]:
            ozet = y.analiz["alinti"] or y.maskeli[:200]
            satirlar.append(f'  <ornek id="{y.id}" duygu="{y.analiz["duygu"]}">{ozet}</ornek>')
        satirlar.append("</tema>")
    satirlar.append("</tema_ozeti>")
    yanit = llm.json_iste(sistem, "\n".join(satirlar), SEMA_BASLIK, max_tokens=8000)
    gecerli = {y.id for y in analizli}
    sonuc, atilan = [], []
    for b in yanit.get("basliklar", [])[:8]:
        ids = [i for i in dict.fromkeys(b.get("yorum_idleri", [])) if i in gecerli]
        if not ids:
            atilan.append(b.get("baslik", "?"))
            continue
        b["yorum_idleri"] = ids
        sonuc.append(b)
    return sonuc, atilan


def tema_tablosu(yorumlu: list[Yanit], gorunur: set[str]) -> dict:
    t = {}
    for ad in TEMALAR:
        lst = [y for y in yorumlu if y.analiz and ad in y.analiz["temalar"]]
        if lst:
            birim = Counter(y.rapor_birimi if y.rapor_birimi in gorunur else "Gizli" for y in lst)
            t[ad] = {"toplam": len(lst), "olumsuz": sum(1 for y in lst if y.analiz["duygu"] in ("olumsuz", "karma")), "birim": birim}
    return dict(sorted(t.items(), key=lambda i: -i[1]["toplam"]))


def anahtar_kelime_riski(y: Yanit) -> list[str]:
    k = katla(y.yorum)
    return [w for w in RISK_KELIMELERI if w in k]


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KONTROL = PatternFill("solid", fgColor="FFF4CE")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, basliklar, genislik=()):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def _isi(hucre, oran):
    if oran is not None:
        hucre.fill = PatternFill("solid", fgColor="FDE2E1" if oran < 0.4 else "FFF4CE" if oran < 0.6 else "E3F4E1")


def rapor_yaz(cikti: Path, s: dict) -> None:
    p = s["puan"]
    wb = Workbook()
    oz = wb.active
    oz.title = "Özet"
    _baslik(oz, ["Boyut", "Yanıt (madde)", "Ortalama (1–5)", "Olumlu (4–5)", "Olumsuz (1–2)"], (28, 13, 13, 12, 12))
    for b in sorted(p["boyutlar"], key=lambda b: p["sirket"][b]["olumlu"] if p["sirket"][b]["olumlu"] is not None else 9):
        x = p["sirket"][b]
        oz.append([b, x["n"], x["ort"], x["olumlu"], x["olumsuz"]])
        oz.cell(oz.max_row, 3).number_format = "0.00"
        for j in (4, 5):
            oz.cell(oz.max_row, j).number_format = "0%"
        _isi(oz.cell(oz.max_row, 4), x["olumlu"])
    e = p["sirket_enps"]
    oz.append([])
    bilgiler = [("Yanıt sayısı", len(s["yanitlar"])), ("Yorum yazan", len(s["yorumlu"])),
                ("eNPS", None if e["enps"] is None else round(e["enps"])), ("Destekçi (9–10) / Nötr (7–8) / Kötüleyen (0–6)",
                                                                           f"{e['destekci']} / {e['notr']} / {e['kotuleyen']}")]
    if s["katilim"]:
        top_c = sum(s["katilim"].values())
        bilgiler.insert(1, ("Katılım oranı", f"%{len(s['yanitlar']) / top_c * 100:.0f} ({len(s['yanitlar'])} / {top_c})"))
    for a, v in bilgiler:
        oz.append([a, v])
    oz.append([])
    for t in (f"Anonimlik: {s['min_yanit']} yanıtın altındaki birimler '{DIGER}' altında birleştirilir; bu grup da küçükse birim tablolarında gösterilmez.",
              "Ters maddeler (ör. 'İş yüküm fazla') 6 − puan olarak çevrilmiştir; tüm boyutlarda yüksek puan olumludur.",
              "Puan analizi koddan, yorum temaları ve sorun başlıkları modelden gelir (mor hücreler). Model: " + s["model"]):
        oz.append([t])

    bb = wb.create_sheet("Birim × Boyut")
    birimler = list(p["birim"])
    _baslik(bb, ["Boyut (olumlu %)"] + birimler + ["Şirket"], [26] + [14] * (len(birimler) + 1))
    for b in p["boyutlar"]:
        bb.append([b] + [p["birim"][u]["boyut"][b]["olumlu"] for u in birimler] + [p["sirket"][b]["olumlu"]])
        for j in range(2, len(birimler) + 3):
            bb.cell(bb.max_row, j).number_format = "0%"
            _isi(bb.cell(bb.max_row, j), bb.cell(bb.max_row, j).value)
    bb.append(["eNPS"] + [None if p["birim"][u]["enps"]["enps"] is None else round(p["birim"][u]["enps"]["enps"]) for u in birimler]
              + [None if e["enps"] is None else round(e["enps"])])
    bb.append(["Yanıt"] + [p["birim"][u]["n"] for u in birimler] + [len(s["yanitlar"])])
    if s["katilim"]:
        bb.append(["Katılım"] + [(p["birim"][u]["n"] / s["katilim"][u]) if s["katilim"].get(u) else None for u in birimler] + [None])
        for j in range(2, len(birimler) + 2):
            bb.cell(bb.max_row, j).number_format = "0%"

    sq = wb.create_sheet("Sorular")
    _baslik(sq, ["Kod", "Soru", "Boyut", "Ters", "Yanıt", "Ortalama / eNPS", "Olumlu", "Olumsuz"], (6, 60, 20, 6, 7, 13, 9, 9))
    for x in p["soru"]:
        q = x["soru"]
        if q.enps:
            sq.append([q.kod, q.metin, q.boyut, "", x["n"], None if x["enps"] is None else round(x["enps"]), None, None])
        else:
            sq.append([q.kod, q.metin, q.boyut, "Evet" if q.ters else "", x["n"], x["ort"], x["olumlu"], x["olumsuz"]])
            sq.cell(sq.max_row, 6).number_format = "0.00"
            for j in (7, 8):
                sq.cell(sq.max_row, j).number_format = "0%"
            _isi(sq.cell(sq.max_row, 7), x["olumlu"])

    sb = wb.create_sheet("Sorun Başlıkları")
    _baslik(sb, ["Başlık", "Açıklama", "Temalar", "Dayanak Yorum", "Örnek Alıntılar (doğrulanmış)", "Olası Aksiyon (öneri)", "Sorumlu / Karar"],
            (30, 60, 26, 10, 60, 40, 22))
    yid = {y.id: y for y in s["yorumlu"]}
    for b in s["basliklar"]:
        alintilar = [f"“{yid[i].analiz['alinti']}”" for i in b["yorum_idleri"] if yid[i].analiz.get("alinti")][:4]
        sb.append([b["baslik"], b["aciklama"], ", ".join(b["temalar"]), len(b["yorum_idleri"]), "\n".join(alintilar), b["olasi_aksiyon"], ""])
        for j in (1, 2, 3, 5, 6):
            sb.cell(sb.max_row, j).fill = AI_DOLGU
        sb.cell(sb.max_row, 7).fill = KONTROL
        for h in sb[sb.max_row]:
            h.alignment = UST

    tm = wb.create_sheet("Temalar")
    gorunur = sorted({u for x in s["temalar"].values() for u in x["birim"]})
    _baslik(tm, ["Tema", "Yorum", "Olumsuz / Karma"] + gorunur, [26, 8, 13] + [12] * len(gorunur))
    for t, x in s["temalar"].items():
        tm.append([t, x["toplam"], x["olumsuz"]] + [x["birim"].get(u, 0) for u in gorunur])

    yr = wb.create_sheet("Yorumlar")
    _baslik(yr, ["Yanıt No", "Birim", "Yorum (maskeli)", "Temalar", "Duygu", "Öneri", "Alıntı", "Hassas", "Kontrol"], (9, 16, 60, 28, 9, 34, 40, 18, 28))
    for y in s["yorumlu"]:
        a = y.analiz
        yr.append([y.id, y.rapor_birimi if y.rapor_birimi in p["birim"] else "Gizli", y.maskeli, ", ".join(a.get("temalar", [])), a.get("duygu", ""),
                   a.get("oneri", ""), a.get("alinti", ""), y.hassas, "; ".join(y.kontrol)])
        for j in (4, 5, 6, 7):
            yr.cell(yr.max_row, j).fill = AI_DOLGU
        if y.hassas:
            yr.cell(yr.max_row, 8).fill = PatternFill("solid", fgColor="FDE2E1")
        for h in yr[yr.max_row]:
            h.alignment = UST
    yr.freeze_panes = "B2"

    hs = wb.create_sheet("Hassas Yorumlar", 1)
    _baslik(hs, ["Yanıt No", "Birim", "Yorum (maskeli)", "Model", "Anahtar Kelime", "Yapılacak / Sorumlu"], (9, 16, 70, 26, 22, 30))
    for y in s["yorumlu"]:
        if y.hassas:
            hs.append([y.id, y.rapor_birimi if y.rapor_birimi in p["birim"] else "Gizli", y.maskeli,
                       y.analiz.get("risk_turu", "") if y.analiz.get("risk_turu") != "yok" else "", ", ".join(anahtar_kelime_riski(y)), ""])
            hs.cell(hs.max_row, 6).fill = KONTROL
            for h in hs[hs.max_row]:
                h.alignment = UST
    hs.append([])
    hs.append(["Not: Anket anonimdir. Bu yorumlar kişiyi bulmaya çalışmadan, ilgili sürece (etik kurul, İSG kurulu, İK) yönlendirilmelidir."])

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Kim", "Açıklama"], (9, 26, 16, 90))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["kim"], u["aciklama"]])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(anket: Path, sorular_yolu: Path, cikti: Path, *, min_yanit: int = 5, katilim: dict[str, int] | None = None, yorum: bool = True,
             evet: bool = False) -> dict:
    sorular = oku_sorular(sorular_yolu)
    yanitlar, uyarilar = oku_anket(anket, sorular)
    if not yanitlar:
        raise llm.LLMHatasi(f"{anket.name}: yanıt bulunamadı.")
    uyarilar += anonimlestir(yanitlar, min_yanit)
    puan = puan_analizi(yanitlar, sorular, min_yanit)
    uyarilar += puan["uyarilar"]
    katilim = {k: v for k, v in (katilim or {}).items()}
    for k in list(katilim):
        if k not in {y.birim for y in yanitlar}:
            uyarilar.append({"onem": "Bilgi", "tur": "Katılım", "kim": k, "aciklama": "Çalışan sayısı verilen birimden yanıt yok"})
    if katilim:   # küçük birimler birleştirildiyse çalışan sayıları da birleştirilir
        birlesik = defaultdict(int)
        for y in {(y.birim, y.rapor_birimi) for y in yanitlar}:
            birlesik[y[1]] += katilim.get(y[0], 0)
        katilim = {k: v for k, v in birlesik.items() if v}
    yorumlu = [y for y in yanitlar if y.yorum]
    for y in yorumlu:
        y.maskeli = maskele(y.yorum)
    basliklar, model = [], "kullanılmadı (--yorum-yok)"
    print(f"[OK] {len(yanitlar)} yanıt · {len(puan['birim'])} birim tablosu · {len(yorumlu)} yorum")
    if yorum and yorumlu:
        llm.onay_al(f"{len(yorumlu)} açık uçlu yorum (e-posta, telefon, TCKN ve unvanlı adlar maskeli; birim bilgisi gönderilmez) analiz için "
                    "gönderilecek.", evet)
        yorumlari_analiz_et(yorumlu)
        temalar = tema_tablosu(yorumlu, set(puan["birim"]))
        basliklar, atilan = sorun_basliklari(yorumlu, temalar)
        for b in atilan:
            uyarilar.append({"onem": "Bilgi", "tur": "Başlık çıkarıldı", "kim": "", "aciklama": f"'{b}' geçerli bir yorum kimliğine dayanmadığı için rapora alınmadı"})
        model = llm.kullanim_ozeti()
    else:
        temalar = {}
    for y in yorumlu:
        r = y.analiz.get("risk_turu", "yok") if y.analiz else "yok"
        kel = anahtar_kelime_riski(y)
        y.hassas = r if r != "yok" else ("Anahtar kelime: " + ", ".join(kel) if kel else "")
        if y.hassas:
            uyarilar.append({"onem": "Yüksek", "tur": "Hassas yorum", "kim": y.id, "aciklama": f"{y.hassas} — ilgili sürece yönlendirin (Hassas Yorumlar)"})
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    uyarilar.sort(key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))
    s = {"yanitlar": yanitlar, "yorumlu": yorumlu, "puan": puan, "temalar": temalar, "basliklar": basliklar, "uyarilar": uyarilar,
         "min_yanit": min_yanit, "katilim": katilim, "model": model}
    rapor_yaz(cikti, s)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Çalışan bağlılık anketi puanlarını birim bazında analiz eder, açık uçlu yorumlardan sorun başlıkları çıkarır.")
    p.add_argument("--anket", type=Path, default=ORNEK / "anket.csv", help="Yanıt No, Birim, [Kıdem], soru kodları (S01…), Görüş ve Öneriler")
    p.add_argument("--sorular", type=Path, default=ORNEK / "sorular.csv", help="Soru Kodu, Soru, Boyut, Ters Madde (Evet/Hayır); eNPS boyutu 0–10")
    p.add_argument("--min-yanit", type=int, default=5, help="Birim tablosunda gösterilecek en az yanıt (anonimlik; varsayılan 5)")
    p.add_argument("--calisan-sayisi", nargs="*", default=[], metavar="BİRİM=SAYI", help="Katılım oranı için birim çalışan sayıları")
    p.add_argument("--yorum-yok", action="store_true", help="Yorumları modele gönderme; yalnız puan analizi")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "baglilik_anketi_analizi.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    for y in (a.anket, a.sorular):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    katilim = {}
    for x in a.calisan_sayisi:
        ad, _, v = x.partition("=")
        if not v.strip().isdigit():
            print(f"[X] --calisan-sayisi 'Birim=sayı' biçiminde olmalı: {x}")
            return 2
        katilim[ad.strip()] = int(v)
    try:
        s = calistir(a.anket, a.sorular, a.cikti, min_yanit=a.min_yanit, katilim=katilim, yorum=not a.yorum_yok, evet=a.evet)
    except (llm.LLMHatasi, ValueError) as h:
        print(f"[X] {h}")
        return 1
    p_ = s["puan"]
    e = p_["sirket_enps"]["enps"]
    print(f"[OK] eNPS {'—' if e is None else f'{e:+.0f}'} · en düşük boyutlar: "
          + ", ".join(f"{b} %{p_['sirket'][b]['olumlu'] * 100:.0f}" for b in sorted(p_["boyutlar"], key=lambda b: p_["sirket"][b]["olumlu"] or 0)[:3]))
    for b in s["basliklar"]:
        print(f"     • {b['baslik']} ({len(b['yorum_idleri'])} yorum)")
    for u in s["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['kim']} {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    if not a.yorum_yok and s["yorumlu"]:
        print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
