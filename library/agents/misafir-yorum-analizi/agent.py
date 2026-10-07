"""
Misafir Yorum Analizi — Workers / Workless AI Agent
Otel › Misafir İlişkileri

1. Kod, yorumları okur; platform puanlarını 10'luk ölçeğe çevirir (Booking 10, Google ve TripAdvisor 5 —
   --olcek ile değiştirilir), haftalık ortalamaları ve platform dağılımını hesaplar.
2. Model her yorumu (çok dilli) Türkçe özetler, konu bazında duygu çıkarır (temizlik, oda, yemek, personel,
   konum, fiyat/değer, havuz-plaj, gürültü, Wi-Fi, check-in/out...), acil durumları (hijyen, sağlık,
   güvenlik, ayrımcılık, kaba davranış) işaretler ve yorumun kendi dilinde cevap taslağı yazar.
3. Kod sonuçları toplar: konu bazında bahsedilme, olumlu/olumsuz dağılımı ve net duygu skoru, önceki haftaya
   göre değişim, puanla duygunun çeliştiği yorumlar ve acil listesi. Hiçbir cevap otomatik yayımlanmaz.

Kullanım:
    python agent.py                                           # örnek yorumlarla dener
    python agent.py --girdi yorumlar.xlsx
    python agent.py --girdi yorumlar.xlsx --olcek Booking=10 Google=5 Tatilsepeti=10 --cevap-yok
"""
from __future__ import annotations

import argparse
import csv
import sys
from collections import Counter, OrderedDict, defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import llm

BURASI = Path(__file__).resolve().parent
PAKET = 25
KONULAR = ["Temizlik", "Oda", "Yemek ve içecek", "Personel ve hizmet", "Konum", "Fiyat / değer", "Havuz / plaj / tesis",
           "Gürültü", "Wi-Fi ve teknoloji", "Check-in / check-out", "Animasyon ve aktiviteler", "Genel deneyim", "Diğer"]
VARSAYILAN_OLCEK = {"booking": 10, "booking.com": 10, "google": 5, "tripadvisor": 5, "holidaycheck": 6, "expedia": 5, "hotels.com": 10}


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def sayi(x) -> float | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return float(x)
    try:
        return float(Decimal(str(x).strip().replace(",", ".")))
    except InvalidOperation:
        return None


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%d.%m.%Y %H:%M"):
        try:
            return datetime.strptime(str(x).strip(), f).date()
        except ValueError:
            pass
    return None


def yorumlari_oku(yol: Path, olcek: dict[str, float]) -> tuple[list[dict], list[str]]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    s = [r for r in s if any(c not in (None, "") for c in r)]
    b = [kucuk(x) for x in s[0]]
    bul = lambda *a: next((i for i, x in enumerate(b) if x in a), None)  # noqa: E731
    k = {"id": bul("yorum no", "id", "no"), "tarih": bul("tarih", "yorum tarihi"), "platform": bul("platform", "kaynak", "site"),
         "puan": bul("puan", "skor", "rating"), "baslik": bul("başlık", "baslik", "title"), "yorum": bul("yorum", "metin", "review", "içerik")}
    if k["yorum"] is None:
        raise SystemExit(f"'Yorum' sütunu bulunamadı. Başlıklar: {s[0]}")
    al = lambda r, a: r[k[a]] if k[a] is not None and k[a] < len(r) else None  # noqa: E731
    yorumlar, uyarilar = [], []
    for i, r in enumerate(s[1:], 1):
        metin = str(al(r, "yorum") or "").strip()
        baslik = str(al(r, "baslik") or "").strip()
        if not metin and not baslik:
            continue
        platform = str(al(r, "platform") or "").strip()
        puan = sayi(al(r, "puan"))
        ust = olcek.get(kucuk(platform))
        if puan is not None and ust is None:
            ust = 5 if puan <= 5 else 10
            uyarilar.append(f"{platform or '(platform yok)'} ölçeği bilinmiyor; {ust:g} üzerinden kabul edildi (--olcek ile verin)")
        yorumlar.append({"id": str(al(r, "id") or f"Y{i}"), "tarih": tarih(al(r, "tarih")), "platform": platform, "puan": puan,
                         "puan10": None if puan is None else puan / ust * 10, "baslik": baslik, "yorum": metin})
    return yorumlar, list(dict.fromkeys(uyarilar))


SEMA = {
    "type": "object",
    "properties": {
        "yorumlar": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "id": {"type": "string"},
                    "dil": {"type": "string"},
                    "ozet_tr": {"type": "string"},
                    "konular": {"type": "array", "items": {
                        "type": "object",
                        "properties": {"konu": {"type": "string", "enum": KONULAR}, "duygu": {"type": "string", "enum": ["olumlu", "olumsuz", "notr"]},
                                       "alinti_tr": {"type": "string"}},
                        "required": ["konu", "duygu", "alinti_tr"], "additionalProperties": False}},
                    "acil": {"type": "boolean"},
                    "acil_neden": {"type": "string"},
                    "cevap_taslagi": {"type": "string"},
                },
                "required": ["id", "dil", "ozet_tr", "konular", "acil", "acil_neden", "cevap_taslagi"],
                "additionalProperties": False,
            },
        },
    },
    "required": ["yorumlar"],
    "additionalProperties": False,
}


def analiz_et(yorumlar: list[dict]) -> dict[str, dict]:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    sonuc = {}
    for i in range(0, len(yorumlar), PAKET):
        parca = yorumlar[i:i + PAKET]
        mesaj = "\n".join(["<konu_listesi>" + "; ".join(KONULAR) + "</konu_listesi>", "<yorumlar>",
                           *[f'<yorum id="{y["id"]}" platform="{y["platform"]}" puan10="{"" if y["puan10"] is None else round(y["puan10"], 1)}">\n'
                             f'{llm.maskele(y["baslik"])}\n{llm.maskele(y["yorum"])}\n</yorum>' for y in parca], "</yorumlar>"])
        yanit = llm.json_iste(sistem, mesaj, SEMA)
        gecerli = {y["id"] for y in parca}
        for s in yanit.get("yorumlar", []):
            if s.get("id") in gecerli:
                sonuc[s["id"]] = s
    return sonuc


def tutarsizlik(y: dict, s: dict) -> str:
    if y["puan10"] is None or not s["konular"]:
        return ""
    olumlu = sum(1 for k in s["konular"] if k["duygu"] == "olumlu")
    olumsuz = sum(1 for k in s["konular"] if k["duygu"] == "olumsuz")
    if y["puan10"] <= 4 and olumsuz == 0 and olumlu > 0:
        return "Düşük puan ama yorum olumlu (yanlış puanlama olabilir)"
    if y["puan10"] >= 9 and olumsuz > olumlu:
        return "Yüksek puan ama yorum ağırlıklı olumsuz"
    return ""


def hafta_basi(t: date) -> date:
    return t - timedelta(days=t.weekday())


def ozetle(yorumlar: list[dict], sonuc: dict[str, dict]) -> dict:
    konu = defaultdict(lambda: Counter())
    haftalik_konu = defaultdict(lambda: defaultdict(Counter))
    for y in yorumlar:
        s = sonuc.get(y["id"])
        if not s:
            continue
        for k in s["konular"]:
            konu[k["konu"]][k["duygu"]] += 1
            if y["tarih"]:
                haftalik_konu[hafta_basi(y["tarih"])][k["konu"]][k["duygu"]] += 1
    tablo = OrderedDict()
    for ad in KONULAR:
        c = konu.get(ad)
        if not c:
            continue
        n = sum(c.values())
        tablo[ad] = {"bahsedilme": n, "olumlu": c["olumlu"], "olumsuz": c["olumsuz"], "notr": c["notr"],
                     "net": (c["olumlu"] - c["olumsuz"]) / n * 100}
    haftalar = sorted({hafta_basi(y["tarih"]) for y in yorumlar if y["tarih"]})
    haftalik = OrderedDict()
    for h in haftalar:
        ys = [y for y in yorumlar if y["tarih"] and hafta_basi(y["tarih"]) == h]
        puanlar = [y["puan10"] for y in ys if y["puan10"] is not None]
        haftalik[h] = {"adet": len(ys), "ort": sum(puanlar) / len(puanlar) if puanlar else None,
                       "olumsuz": sum(c["olumsuz"] for c in haftalik_konu[h].values())}
    if len(haftalar) >= 2:
        son, onceki = haftalar[-1], haftalar[-2]
        for ad, x in tablo.items():
            a, b = haftalik_konu[son].get(ad, Counter()), haftalik_konu[onceki].get(ad, Counter())
            x["olumsuz_degisim"] = a["olumsuz"] - b["olumsuz"]
    return {"konu": tablo, "haftalik": haftalik}


BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
ONAY_DOLGU = PatternFill("solid", fgColor="FFF4CE")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def rapor_yaz(cikti: Path, yorumlar: list[dict], sonuc: dict, ozet: dict, uyarilar: list[str], cevap: bool) -> None:
    wb = Workbook()
    k = wb.active
    k.title = "Konu Analizi"
    _baslik(k, ["Konu", "Bahsedilme", "Olumlu", "Olumsuz", "Nötr", "Net Duygu (−100…+100)", "Olumsuz Değişim (son hafta)"], (24, 12, 9, 9, 8, 20, 22))
    for ad, x in sorted(ozet["konu"].items(), key=lambda i: i[1]["net"]):
        k.append([ad, x["bahsedilme"], x["olumlu"], x["olumsuz"], x["notr"], round(x["net"]), x.get("olumsuz_degisim")])
        k.cell(k.max_row, 6).fill = KIRMIZI if x["net"] < 0 else YESIL
    n = len(ozet["konu"])
    if n:
        g = BarChart()
        g.type, g.grouping, g.overlap = "bar", "stacked", 100
        g.title, g.height, g.width = "Konu bazında olumlu / olumsuz", 9, 18
        g.add_data(Reference(k, min_col=3, max_col=4, min_row=1, max_row=1 + n), titles_from_data=True)
        g.set_categories(Reference(k, min_col=1, min_row=2, max_row=1 + n))
        k.add_chart(g, "I2")
    k.append([])
    k.append(["Haftalık", "Yorum", "Ort. Puan (10)", "Olumsuz Bahsedilme"])
    for h in k[k.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for h_, x in ozet["haftalik"].items():
        k.append([f"{h_:%d.%m.%Y} haftası", x["adet"], None if x["ort"] is None else round(x["ort"], 2), x["olumsuz"]])
    k.append([])
    for u in uyarilar:
        k.append(["Uyarı", u])
    k.append(["Model", llm.kullanim_ozeti()])
    k.append(["Not", "Konu ve duygu modelden, puan ölçekleme, sayımlar ve tutarsızlık kontrolü koddan gelir. Cevaplar otomatik yayımlanmaz."])

    a = wb.create_sheet("Acil", 0)
    _baslik(a, ["Yorum No", "Tarih", "Platform", "Puan (10)", "Neden", "Özet", "Yapılacak / Sorumlu"], (10, 12, 12, 10, 40, 60, 30))
    for y in yorumlar:
        s = sonuc.get(y["id"])
        if s and s["acil"]:
            a.append([y["id"], y["tarih"], y["platform"], None if y["puan10"] is None else round(y["puan10"], 1), s["acil_neden"], s["ozet_tr"], ""])
            a.cell(a.max_row, 2).number_format = "DD.MM.YYYY"
            a.cell(a.max_row, 5).fill = KIRMIZI
            a.cell(a.max_row, 7).fill = ONAY_DOLGU
            for h in a[a.max_row]:
                h.alignment = UST

    d = wb.create_sheet("Yorumlar")
    bas = ["Yorum No", "Tarih", "Platform", "Puan", "Puan (10)", "Dil", "Özet (TR)", "Olumlu Konular", "Olumsuz Konular", "Alıntılar", "Acil",
           "Kontrol"] + (["Cevap Taslağı", "Onay"] if cevap else [])
    _baslik(d, bas, (10, 12, 12, 7, 9, 5, 50, 28, 28, 60, 7, 30) + ((60, 10) if cevap else ()))
    for y in yorumlar:
        s = sonuc.get(y["id"])
        if s:
            olumlu = ", ".join(k_["konu"] for k_ in s["konular"] if k_["duygu"] == "olumlu")
            olumsuz = ", ".join(k_["konu"] for k_ in s["konular"] if k_["duygu"] == "olumsuz")
            satir = [y["id"], y["tarih"], y["platform"], y["puan"], None if y["puan10"] is None else round(y["puan10"], 1), s["dil"], s["ozet_tr"],
                     olumlu, olumsuz, "\n".join(f"{k_['konu']}: {k_['alinti_tr']}" for k_ in s["konular"]), "EVET" if s["acil"] else "",
                     tutarsizlik(y, s)] + ([s["cevap_taslagi"], ""] if cevap else [])
        else:
            satir = [y["id"], y["tarih"], y["platform"], y["puan"], y["puan10"], "", "AI yanıt vermedi"] + [""] * (5 + (2 if cevap else 0))
        d.append(satir)
        d.cell(d.max_row, 2).number_format = "DD.MM.YYYY"
        for c in (6, 7, 8, 9, 10) + ((13,) if cevap else ()):
            d.cell(d.max_row, c).fill = AI_DOLGU
        if cevap:
            d.cell(d.max_row, 14).fill = ONAY_DOLGU
        if satir[11]:
            d.cell(d.max_row, 12).fill = KIRMIZI
        for h in d[d.max_row]:
            h.alignment = UST
    d.freeze_panes = "B2"
    d.auto_filter.ref = d.dimensions
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(girdi: Path, cikti: Path, olcek: dict[str, float] | None = None, cevap: bool = True, evet: bool = False) -> dict:
    o = dict(VARSAYILAN_OLCEK)
    o.update({kucuk(k): v for k, v in (olcek or {}).items()})
    yorumlar, uyarilar = yorumlari_oku(girdi, o)
    if not yorumlar:
        raise llm.LLMHatasi(f"{girdi.name}: yorum bulunamadı.")
    print(f"[OK] {len(yorumlar)} yorum · platformlar: " + ", ".join(f"{p} {n}" for p, n in Counter(y["platform"] for y in yorumlar).most_common()))
    llm.onay_al(f"{len(yorumlar)} yorumun başlık ve metni (e-posta/telefon maskeli) analiz için gönderilecek.", evet)
    sonuc = analiz_et(yorumlar)
    ozet = ozetle(yorumlar, sonuc)
    rapor_yaz(cikti, yorumlar, sonuc, ozet, uyarilar, cevap)
    return {"yorumlar": yorumlar, "sonuc": sonuc, "ozet": ozet, "uyarilar": uyarilar}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Misafir yorumlarını konu ve duygu bazında analiz eder, cevap taslağı hazırlar.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "yorumlar.csv",
                   help="Yorumlar (.xlsx/.csv): Yorum; Yorum No, Tarih, Platform, Puan, Başlık isteğe bağlı")
    p.add_argument("--olcek", nargs="*", default=[], metavar="PLATFORM=EN_YÜKSEK_PUAN", help="Platform puan ölçekleri, ör. Booking=10 Google=5")
    p.add_argument("--cevap-yok", action="store_true", help="Cevap taslağı üretme")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "misafir_yorumlari.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    olcek = {}
    for x in a.olcek:
        ad, _, v = x.partition("=")
        if not v or sayi(v) is None:
            print(f"[X] --olcek 'Platform=puan' biçiminde olmalı: {x}")
            return 2
        olcek[ad] = sayi(v)
    try:
        s = calistir(a.girdi, a.cikti, olcek, not a.cevap_yok, a.evet)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    acil = [i for i, x in s["sonuc"].items() if x["acil"]]
    en_kotu = sorted(s["ozet"]["konu"].items(), key=lambda i: i[1]["net"])[:3]
    print(f"[OK] {len(s['sonuc'])}/{len(s['yorumlar'])} yorum analiz edildi · acil: {len(acil)}")
    for ad, x in en_kotu:
        print(f"     {ad}: net duygu {x['net']:+.0f} ({x['olumsuz']} olumsuz / {x['bahsedilme']} bahsedilme)")
    for u in s["uyarilar"]:
        print(f"[!] {u}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
