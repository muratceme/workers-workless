"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri" / "yorumlar.csv"
MESAJLAR = []
YANIT = {  # id → (konular, acil)
    "Y001": ([("Konum", "olumlu"), ("Personel ve hizmet", "olumlu"), ("Yemek ve içecek", "olumlu")], False),
    "Y002": ([("Temizlik", "olumsuz")], False),
    "Y003": ([("Havuz / plaj / tesis", "olumlu"), ("Oda", "olumsuz"), ("Wi-Fi ve teknoloji", "olumsuz")], False),
    "Y004": ([("Temizlik", "olumlu"), ("Gürültü", "olumsuz")], False),
    "Y005": ([("Genel deneyim", "olumlu")], False),
    "Y006": ([("Yemek ve içecek", "olumsuz")], True),
    "Y007": ([("Personel ve hizmet", "olumlu"), ("Temizlik", "olumlu")], False),
    "Y008": ([("Fiyat / değer", "olumlu"), ("Gürültü", "olumsuz"), ("Check-in / check-out", "olumsuz")], False),
    "Y009": ([("Personel ve hizmet", "olumsuz")], True),
    "Y010": ([("Temizlik", "olumlu"), ("Gürültü", "olumsuz")], False),
}


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    idler = re.findall(r'<yorum id="([^"]+)"', kullanici)
    return {"yorumlar": [{"id": i, "dil": "tr", "ozet_tr": "özet", "konular": [{"konu": k, "duygu": d, "alinti_tr": "..."} for k, d in YANIT[i][0]],
                          "acil": YANIT[i][1], "acil_neden": "Hijyen" if YANIT[i][1] else "", "cevap_taslagi": "Merhaba"} for i in idler]
            + [{"id": "Y999", "dil": "tr", "ozet_tr": "", "konular": [], "acil": True, "acil_neden": "uydurma", "cevap_taslagi": ""}]}


class OkumaTesti(unittest.TestCase):
    def test_olcek(self):
        y, u = agent.yorumlari_oku(ORNEK, agent.VARSAYILAN_OLCEK)
        p = {x["id"]: x["puan10"] for x in y}
        self.assertEqual((p["Y001"], p["Y002"], p["Y003"]), (9.2, 4.0, 8.0))     # Booking 10, Google 5, TripAdvisor 5
        self.assertEqual(u, [])

    def test_bilinmeyen_platform(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp) / "y.csv"
            t.write_text("Platform;Puan;Yorum\nX;4;iyi\nZ;8;güzel\n", encoding="utf-8")
            y, u = agent.yorumlari_oku(t, {})
        self.assertEqual([x["puan10"] for x in y], [8.0, 8.0])
        self.assertEqual(len(u), 2)

    def test_tutarsizlik(self):
        self.assertIn("Düşük puan", agent.tutarsizlik({"puan10": 2}, {"konular": [{"duygu": "olumlu"}]}))
        self.assertEqual(agent.tutarsizlik({"puan10": 2}, {"konular": [{"duygu": "olumsuz"}]}), "")


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_uctan_uca(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "m.xlsx"
            s = agent.calistir(ORNEK, cikti, evet=True)
            self.assertNotIn("Y999", s["sonuc"])
            k = s["ozet"]["konu"]
            self.assertEqual((k["Gürültü"]["bahsedilme"], k["Gürültü"]["olumsuz"], round(k["Gürültü"]["net"])), (3, 3, -100))
            self.assertEqual((k["Temizlik"]["olumlu"], k["Temizlik"]["olumsuz"]), (3, 1))
            self.assertEqual(round(k["Temizlik"]["net"]), 50)
            self.assertEqual(list(s["ozet"]["haftalik"])[0].isoformat(), "2026-09-21")
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Acil", "Konu Analizi", "Yorumlar"])
            self.assertEqual(wb["Acil"].max_row, 3)                      # Y006, Y009
            ws = wb["Yorumlar"]
            self.assertEqual(ws.cell(1, ws.max_column).value, "Onay")

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_cevapsiz(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "m.xlsx"
            agent.calistir(ORNEK, cikti, cevap=False, evet=True)
            ws = load_workbook(cikti)["Yorumlar"]
            self.assertEqual(ws.cell(1, ws.max_column).value, "Kontrol")


if __name__ == "__main__":
    unittest.main()
