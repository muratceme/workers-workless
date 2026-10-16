"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import json
import os
import re
import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
CEVAP = json.loads((Path(__file__).parent / "sahte_cevaplar.json").read_text(encoding="utf-8"))
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    """Stok-002'de ilk turda değişen ön kaputu ekspertiz özetine yazmaz."""
    MESAJLAR.append(kullanici)
    no = re.search(r'<arac no="([^"]+)"', kullanici)[1]
    return copy.deepcopy(CEVAP[f"{no}_hatali" if no == "Stok-002" and "<duzeltme>" not in kullanici else no])


def hazir():
    return agent.araclari_oku(ORNEK / "araclar.csv"), agent.ekspertiz_oku(ORNEK / "ekspertiz_notlari.txt")


class KodTesti(unittest.TestCase):
    def test_okuma(self):
        k, e = hazir()
        self.assertEqual((k["Stok-002"]["km"], k["Stok-002"]["fiyat"], k["Stok-002"]["renk"]), ("150000", "760.000", "Gri"))
        self.assertEqual((e["Stok-001"]["boyali"], e["Stok-001"]["lokal"], e["Stok-001"]["tramer"]), ([], ["sağ arka çamurluk"], None))
        self.assertEqual((e["Stok-002"]["degisen"], e["Stok-002"]["tramer"]), (["ön kaput"], D(45000)))
        self.assertEqual(e["Stok-001"]["not"], "Tek kullanıcılı, sigara içilmemiş.")
        self.assertEqual(agent.parcalar("sol ön kapı ve bagaj kapağı; tavan"), ["sol ön kapı", "bagaj kapağı", "tavan"])

    def test_denetim(self):
        k, e = hazir()
        d = lambda n, c, s=60: agent.denetle(k[n], e[n], c, s)
        self.assertEqual([d("Stok-001", CEVAP["Stok-001"]), d("Stok-002", CEVAP["Stok-002"])], [[], []])
        self.assertEqual(d("Stok-002", CEVAP["Stok-002_hatali"]), ["Stok-002: değişen parça ekspertiz özetinde yok: ön kaput"])
        c = copy.deepcopy(CEVAP["Stok-002"])
        c["baslik"] = "Örnek Marka Model A · Kazasız · Hatasız"
        c["aciklama"] = c["aciklama"].replace("150.000", "140.000") + " Değişen parça yoktur. Sunroof mevcut."
        c["donanim"].append("sunroof")
        c["ekspertiz_ozeti"] = "Sol ön çamurluk boyalıdır. Ön kaput değişmiştir."
        s = d("Stok-002", c, 30)
        for p in ("başlık 39 karakter", "model yılı (2020)", "km (150000)", "olmayan sayı: 140.000", "tramer (hasar kaydı) tutarı 45.000",
                  "'hatasiz'", "'degisen parca yoktur'", "'kazasiz'", "notlarda olmayan donanım: sunroof"):
            self.assertTrue(any(p in x for x in s), p)
        c = copy.deepcopy(CEVAP["Stok-001"])
        c["ekspertiz_ozeti"] = "Araç boyasızdır, değişen ve tramer kaydı yoktur."
        s = d("Stok-001", c)
        self.assertTrue(any("lokal boyalı parça ekspertiz özetinde yok" in x for x in s) and any("boyalı parça varken 'boyasizdir'" in x for x in s))
        self.assertFalse(any("değişen parça varken" in x or "tramer / değişen varken" in x for x in s))   # değişen ve tramer yok: serbest


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ornek_ve_duzeltme(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "i.xlsx"
            self.assertEqual(agent.main(["--evet", "--cikti", str(cikti), "--ek-satir", "Yetki Belgesi No: 0000"]), 0)
            self.assertEqual(len(MESAJLAR), 3)
            ws = load_workbook(cikti)["İlanlar"]
            self.assertEqual([ws.cell(r, 11).value for r in (2, 3)], ["Geçti", "Geçti"])
            self.assertIn("Ön kaput değişmiştir", ws.cell(3, 10).value)
            t = cikti.with_suffix(".txt").read_text(encoding="utf-8")
            self.assertEqual(t.count("Yetki Belgesi No: 0000"), 2)
            self.assertIn("Ekspertiz:\nSol ön çamurluk boyalıdır.", t)

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_secim_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(ORNEK / "araclar.csv", ORNEK / "ekspertiz_notlari.txt", Path(tmp) / "i.xlsx", secili=["Stok-001"], evet=True)
            self.assertEqual([a["no"] for a in s["araclar"]], ["Stok-001"])
            y = Path(tmp) / "e.txt"
            y.write_text("Araç: Stok-001\nBoyalı: yok\n", encoding="utf-8")
            self.assertEqual(agent.main(["--ekspertiz", str(y), "--evet"]), 1)            # Stok-002'nin notu yok


if __name__ == "__main__":
    unittest.main()
