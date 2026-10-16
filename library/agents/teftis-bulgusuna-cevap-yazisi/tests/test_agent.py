"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import json
import os
import re
import sys
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
CEVAP = json.loads((Path(__file__).parent / "sahte_cevaplar.json").read_text(encoding="utf-8"))
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    """T-26-14'te ilk turda BT değişikliği planlanırken bulguyu 'Tamamlandı' sayar ve planı atlar."""
    MESAJLAR.append(kullanici)
    no = re.search(r'<bulgu no="([^"]+)"', kullanici)[1]
    return copy.deepcopy(CEVAP[f"{no}_hatali" if no == "T-26-14" and "<duzeltme>" not in kullanici else no])


def bulgular():
    _, b = agent.bulgulari_oku(ORNEK / "bulgular.csv")
    n = agent.notlari_oku(ORNEK / "sube_notlari.txt", {})
    return {x["no"]: x | {"not": n[x["no"]]} for x in b}


class KodTesti(unittest.TestCase):
    def test_okuma(self):
        b = bulgular()
        self.assertEqual(list(b), ["T-26-14", "T-26-15", "T-26-16"])
        self.assertEqual(b["T-26-14"]["acilis"], date(2026, 9, 22))
        self.assertTrue(b["T-26-16"]["not"].startswith("Yedek anahtar teslim tutanağı"))
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "b.csv"
            y.write_text("No;Konu\n1;x\n", encoding="utf-8")
            with self.assertRaises(llm.LLMHatasi):
                agent.bulgulari_oku(y)

    def test_denetim(self):
        b = bulgular()
        self.assertEqual([agent.denetle(b[n], CEVAP[n]) for n in b], [[], [], []])
        self.assertEqual(agent.denetle(b["T-26-14"], CEVAP["T-26-14_hatali"]),
                         ["T-26-14: nottaki tarih cevapta yok: 30.11.2026 (bu tarihli işi atlamayın)",
                          "T-26-14: notta henüz yapılmamış iş var ('planliyor'); durum 'Tamamlandı' olamaz ve planlanan aksiyon yazılmalı"])
        c = copy.deepcopy(CEVAP["T-26-16"])
        c.update(kapanis_tarihi="", cevap=c["cevap"] + " 3 personele eğitim verildi.")
        s = agent.denetle(b["T-26-16"], c)
        self.assertTrue(any("kapanış tarihi gerekli" in x for x in s) and any("olmayan sayı: 3" in x for x in s))
        c = copy.deepcopy(CEVAP["T-26-15"]) | {"termin": "", "planlanan_aksiyonlar": ["Kalanlar tamamlanacak"]}
        self.assertTrue(any("termin gerekli" in x for x in agent.denetle(b["T-26-15"], c)))
        c = copy.deepcopy(CEVAP["T-26-15"]) | {"termin": "01.09.2026"}
        self.assertTrue(any("açılışından önce" in x or "notta olmayan tarih" in x for x in agent.denetle(b["T-26-15"], c)))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ornek_ve_duzeltme(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "c.xlsx"
            self.assertEqual(agent.main(["--evet", "--cikti", str(cikti)]), 0)
            self.assertEqual(len(MESAJLAR), 4)                          # 3 bulgu + T-26-14 düzeltmesi
            wb = load_workbook(cikti)
            ws = wb["Bulgular"]
            bas = [c.value for c in ws[1]]
            self.assertEqual(bas[:10], ["Bulgu No", "Bulgu", "Birim", "Önem", "Açılış Tarihi", "Termin", "Kapanış Tarihi", "Aksiyon Sahibi",
                                        "Termin Revizyon Sayısı", "Tekrar"])
            r = {ws.cell(i, 1).value: [c.value for c in ws[i]] for i in range(2, 5)}
            self.assertEqual((r["T-26-14"][5], r["T-26-14"][6], r["T-26-14"][10]), (datetime(2026, 11, 30), None, "Devam ediyor"))
            self.assertEqual((r["T-26-16"][6], r["T-26-16"][10]), (datetime(2026, 9, 25), "Tamamlandı"))
            y = "\n".join(str(c.value or "") for c in wb["Cevap Yazısı"]["A"])
            self.assertIn("Konu: Kadıköy Şubesi teftiş / kontrol bulgularına ilişkin cevaplarımız", y)
            self.assertTrue(y.rstrip().endswith("Saygılarımızla,\nŞube Operasyon Yöneticisi"))

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_notsuz_bulgu_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            n = Path(tmp) / "n.txt"
            n.write_text("Bulgu: T-26-16\nYedek anahtar teslim tutanağı 25.09.2026'da yeniden düzenlenip iki yetkiliye imzalatıldı.\n", encoding="utf-8")
            s = agent.calistir(ORNEK / "bulgular.csv", n, Path(tmp) / "c.xlsx", evet=True)
            self.assertEqual(([b["no"] for b in s["bulgular"]], [b["no"] for b in s["eksik"]]), (["T-26-16"], ["T-26-14", "T-26-15"]))
            self.assertIn("şube notu yok", s["txt"].read_text(encoding="utf-8"))
            self.assertEqual(agent.main(["--notlar", str(Path(tmp) / "yok.txt"), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
