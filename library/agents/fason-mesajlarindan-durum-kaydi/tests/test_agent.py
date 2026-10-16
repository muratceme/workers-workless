"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import json
import os
import sys
import tempfile
import unittest
from datetime import date, datetime
from decimal import Decimal as D
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
CEVAP = json.loads((Path(__file__).parent / "sahte_cevaplar.json").read_text(encoding="utf-8"))
METIN = {"FS-06 yıkamada": "M1", "kalan 900'ü yarın": "M2", "900 adet teslim": "M3", "FS-08 için": "M4", "FS-03 kalan": "M5"}
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    """M2'de ilk turda 'yarın göndereceğiz' sözünü dönüş sayar."""
    MESAJLAR.append(kullanici)
    m = next(v for k, v in METIN.items() if k in kullanici.split("<mesaj", 1)[1].split("</mesaj>")[0])
    return copy.deepcopy(CEVAP["M2_hatali" if m == "M2" and "<duzeltme>" not in kullanici else m])


def hazir():
    return agent.mesajlari_oku(ORNEK / "mesajlar.txt"), agent.isleri_oku(ORNEK / "fason_isler.csv")


class KodTesti(unittest.TestCase):
    def test_okuma(self):
        m, i = hazir()
        self.assertEqual([(x["id"], x["gun"], x["gonderen"]) for x in m][:2], [("M1", date(2026, 10, 7), "Yıkamacı E"), ("M2", date(2026, 10, 7), "Atölye B")])
        self.assertEqual((i["FS-07"]["adet"], i["FS-07"]["termin"], i["FS-01"]["kapandi"]), (D(1500), date(2026, 10, 10), True))
        self.assertEqual(agent.donusleri_oku(ORNEK / "donusler.csv")["FS-03"], D(1800))

    def test_goreli_tarih(self):
        g = date(2026, 10, 8)                                                    # perşembe
        a = lambda s: sorted(agent.aday_tarihler(s, g))
        self.assertEqual(a("yarın gönderiyoruz"), [date(2026, 10, 9)])
        self.assertEqual(a("cumaya yetişir"), [date(2026, 10, 9)])
        self.assertEqual(a("perşembe teslim"), [g, date(2026, 10, 15)])           # aynı gün ya da haftaya
        self.assertEqual(a("ayın 22'sinde"), [date(2026, 10, 22)])
        self.assertEqual(a("ayın 5'inde"), [date(2026, 11, 5)])                  # geçmiş gün → gelecek ay
        self.assertEqual(a("15.10.2026'da ve haftaya"), [date(2026, 10, 15)])
        self.assertEqual(a("2 adet 45 tanesinde"), [])

    def test_denetim(self):
        m, i = hazir()
        b = {x["id"]: x for x in m}
        self.assertEqual([agent.denetle(b[k], i, CEVAP[k]) for k in ("M1", "M2", "M3", "M4", "M5")], [[], [], [], [], []])
        s = agent.denetle(b["M2"], i, CEVAP["M2_hatali"])
        self.assertEqual(len(s), 2)
        self.assertIn("mesaj tarihinden sonra", s[0])
        self.assertIn("tamamlanmış teslim ifadesi yok", s[1])
        c = copy.deepcopy(CEVAP["M1"])
        c["kayitlar"][1].update(adet="1500", tarih="10.10.2026", kalite_sonucu="Kabul")
        c["kayitlar"].append(c["kayitlar"][0] | {"is_no": "FS-99"})
        s = agent.denetle(b["M1"], i, c)
        for p in ("adet '1500' mesajda yazmıyor", "söz tarihi 10.10.2026 mesajdan çıkmıyor", "yalnız dönüşte", "iş no listede yok"):
            self.assertTrue(any(p in x for x in s), p)
        c = copy.deepcopy(CEVAP["M4"])
        c["kayitlar"] = []
        self.assertEqual(agent.denetle(b["M4"], i, c), ["M4: mesajda geçen iş kayda alınmamış: FS-08"])
        c = copy.deepcopy(CEVAP["M3"])
        c["kayitlar"][0].update(hatali_adet="900", adet="45", kanit="teslim edildi")
        self.assertTrue(any("hatalı adet dönen adetten büyük" in x for x in agent.denetle(b["M3"], i, c)))

    def test_hesap(self):
        m, i = hazir()
        s = {x["id"]: CEVAP[x["id"]] for x in m}
        h = agent.hesapla(m, i, agent.donusleri_oku(ORNEK / "donusler.csv"), s)
        self.assertEqual(h["donen"]["FS-07"], D(1500))
        self.assertEqual([(x["is_no"], x["fark"], x["kalan"]) for x in h["sozler"]], [("FS-06", 2, D(2500)), ("FS-07", -2, D(900)), ("FS-08", 2, D(2000)), ("FS-03", 12, D(1200))])
        h = agent.hesapla(m, i, {"FS-07": D(1000)}, s)                         # önceki 1000 + 900 > 1500
        self.assertTrue(any("aşıyor" in " ".join(v) for v in h["notlar"].values()))
        self.assertTrue(any("kalan 500" in " ".join(v) for v in h["notlar"].values()))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ornek_ve_duzeltme(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "f.xlsx"
            self.assertEqual(agent.main(["--evet", "--cikti", str(cikti)]), 0)
            self.assertEqual(len(MESAJLAR), 6)                          # 5 mesaj + M2 düzeltmesi
            self.assertNotIn("FS-01 ·", MESAJLAR[0])                    # kapanmış iş gönderilmez
            d = [[c.value for c in r] for r in load_workbook(Path(tmp) / "f_donusler.xlsx").active.iter_rows()]
            self.assertEqual(d, [["İş No", "Dönüş Tarihi", "Dönen Adet", "Kalite Sonucu", "Hatalı Adet", "Hata Açıklaması"],
                                 ["FS-07", datetime(2026, 10, 8), 900, "Kabul", 45, "Etek boyu ölçü dışı (tamire geri alındı)"]])
            k = [[c.value for c in r] for r in load_workbook(cikti)["Kayıtlar"].iter_rows(min_row=2)]
            self.assertEqual([x[6] for x in k], ["Sorun", "Söz", "Söz", "Dönüş", "Sorun", "Söz", "Söz"])

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(agent.main(["--mesajlar", str(Path(tmp) / "yok.txt"), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
