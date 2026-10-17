"""
Fason Mesajlarından Durum Kaydı — Gerisi Bende AI Agent
Tekstil › Fason Takip Sorumlusu

Fason atölyelerden gelen kısa mesajları (WhatsApp, SMS, e-posta) iş takip kaydına dönüştürür: dönüşler (teslimler),
verilen teslim sözleri ve sorunlar.
1. Kod açık fason işleri ("Fason Atölye İş Takibi" kod bloğunun iş listesi), önceki dönüşleri (isteğe bağlı) ve
   "[GG.AA.YYYY SS:DD] Gönderen: mesaj" satırlarını okur; telefon, e-posta, TCKN ve --gizle adları maskelenir.
2. Model her mesajdan kayıtları çıkarır: tür (Dönüş / Söz / Sorun), iş no, tarih, adet, hatalı adet, kalite, kanıt.
3. Kod denetler: iş no listede ve mesajda var mı; mesajdaki her iş no kayda alınmış mı; adetler mesajda yazıyor mu;
   tarih mesajda yazan ya da mesaj tarihine göre "bugün / yarın / cuma / ayın 14'ü" ifadesinden çıkan tarih mi; dönüş
   gerçekten geçmiş zamanlı bir teslim mi (göndereceğiz / getiririz dönüş değildir) ve mesaj tarihinden sonra mı; kanıt
   mesajda aynen var mı. Sorun varsa yalnız o mesaj modele geri gönderilir.
4. Kod kalan adedi, termin farkını ve çıkan adedi aşan dönüşleri hesaplar; kod bloğunun dönüş dosyasını yazar.

Kullanım:
    python agent.py                                                         # örnek 5 mesaj
    python agent.py --mesajlar whatsapp.txt --isler fason_isler.xlsx --donusler donusler.xlsx
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}
GUNLER = {"pazartesi": 0, "sali": 1, "carsamba": 2, "persembe": 3, "cumartesi": 5, "cuma": 4, "pazar": 6}
TESLIM = re.compile(r"\b(teslim (ettik|edildi|edilmistir|ettim|aldiniz)|gonderdik|gonderildi|gonderdim|getirdik|getirildi|biraktik|ulasti|yola cikti|"
                    r"sevk ettik|sevk edildi)\b")
SATIR = re.compile(r"^\s*\[(\d{1,2}\.\d{1,2}\.\d{4})\s+(\d{1,2})[:.](\d{2})\]\s*([^:]{1,40}):\s*(.+)$")
IS_NO = re.compile(r"\b[A-Z]{1,4}-?\d{1,6}\b")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def sayi(x) -> Decimal | None:
    s = re.sub(r"[^\d,.\-]", "", str(x if x is not None else ""))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1 or re.fullmatch(r"\d{1,3}\.\d{3}", s):
        s = s.replace(".", "")
    try:
        return Decimal(s) if s not in ("", "-", ".") else None
    except InvalidOperation:
        return None


def sayilar(metin: str) -> set[Decimal]:
    temiz = re.sub(r"\b\d{1,2}\.\d{1,2}\.\d{4}\b|\b\d{1,2}:\d{2}\b|" + IS_NO.pattern, " ", str(metin))
    return {v for x in re.findall(r"\d+(?:[.,]\d+)*", temiz) if (v := sayi(x)) is not None}


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    m = re.search(r"\b(\d{1,2})\.(\d{1,2})\.(\d{4})\b", str(x or ""))
    try:
        return date(int(m[3]), int(m[2]), int(m[1])) if m else None
    except ValueError:
        return None


def aday_tarihler(metin: str, gun: date) -> set[date]:
    """Mesajda yazan ya da mesaj tarihine göre göreli ifadeden çıkan tarihler."""
    k = katla(metin)
    out = {t for x in re.findall(r"\b\d{1,2}\.\d{1,2}\.\d{4}\b", metin) if (t := tarih_coz(x))}
    for ifade, fark in (("bugun", 0), ("dun", -1), ("yarin", 1), ("obur gun", 2), ("ertesi gun", 1), ("haftaya", 7)):
        if re.search(rf"\b{ifade}\w*", k):
            out.add(gun + timedelta(days=fark))
    for m in re.finditer(r"\b(pazartesi|sali|carsamba|persembe|cumartesi|cuma|pazar)\w*", k):
        f = (GUNLER[m[1]] - gun.weekday()) % 7
        out.add(gun + timedelta(days=f or 7))
        if f == 0:
            out.add(gun)
    for m in re.finditer(r"\b(?:ayin )?(\d{1,2}) s?[iu]?n?d?[ea]\b|\bayin (\d{1,2})\b", k):
        n = int(m[1] or m[2])
        y, a = (gun.year, gun.month) if n >= gun.day else ((gun.year + 1, 1) if gun.month == 12 else (gun.year, gun.month + 1))
        try:
            out.add(date(y, a, n))
        except ValueError:
            pass
    return out


# ----------------------------------------------------------------------------
# Girdi
# ----------------------------------------------------------------------------

def tablo_oku(yol: Path) -> list[dict]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, read_only=True, data_only=True)
        s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=metin.split("\n")[0].count)))
    s = [r for r in s if any(c not in (None, "") for c in r)]
    bas = [katla(c) for c in s[0]] if s else []
    return [dict(zip(bas, list(r) + [None] * (len(bas) - len(r)))) for r in s[1:]]


def alan(r: dict, *adlar):
    return next((r[a] for a in adlar if r.get(a) not in (None, "")), None)


def isleri_oku(yol: Path) -> dict[str, dict]:
    out = {}
    for r in tablo_oku(yol):
        no = str(alan(r, "is no", "fason no", "is emri no", "no") or "").strip()
        if no:
            out[no] = {"no": no, "atolye": str(alan(r, "atolye", "fasoncu", "tedarikci") or ""), "islem": str(alan(r, "islem", "operasyon", "is turu") or ""),
                       "adet": sayi(alan(r, "cikan adet", "adet", "gonderilen adet")), "termin": tarih_coz(alan(r, "termin", "donus termini", "teslim termini")),
                       "kapandi": katla(alan(r, "kapandi", "kapali", "durum")).startswith(("evet", "kapa"))}
    if not out:
        raise llm.LLMHatasi(f"{yol.name}: 'İş No' sütunu olan iş listesi gerekli.")
    return out


def donusleri_oku(yol: Path | None) -> dict[str, Decimal]:
    t = defaultdict(Decimal)
    if yol:
        for r in tablo_oku(yol):
            no, a = str(alan(r, "is no", "fason no", "is emri no", "no") or "").strip(), sayi(alan(r, "donen adet", "adet", "gelen adet"))
            if no and a:
                t[no] += a
    return t


def takma_adlar(adlar: list[str]) -> dict[str, str]:
    return {a: f"[KİŞİ-{i}]" for i, a in enumerate(dict.fromkeys(x.strip() for x in adlar if x.strip()), 1)}


def maskele(metin: str, kisiler: dict[str, str]) -> str:
    for g, t in sorted(kisiler.items(), key=lambda x: -len(x[0])):
        metin = re.sub(re.escape(g), t, metin, flags=re.I)
    return llm.maskele(metin)


def mesajlari_oku(yol: Path, gizle: list[str] | None = None) -> list[dict]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: mesaj dosyası bulunamadı.")
    kisiler = takma_adlar(gizle or [])
    out = []
    for s in belge.metin_oku(yol).replace("\r", "").splitlines():
        m = SATIR.match(s)
        if m:
            g = tarih_coz(m[1])
            if not g:
                raise llm.LLMHatasi(f"{yol.name}: tarih okunamadı: '{m[1]}'")
            out.append({"id": f"M{len(out) + 1}", "gun": g, "zaman": f"{g:%d.%m.%Y} {int(m[2]):02d}:{m[3]}", "gonderen": maskele(m[4].strip(), kisiler),
                        "metin": maskele(m[5].strip(), kisiler)})
        elif out and s.strip():
            out[-1]["metin"] += " " + maskele(s.strip(), kisiler)
    if not out:
        raise llm.LLMHatasi(f"{yol.name}: '[GG.AA.YYYY SS:DD] Gönderen: mesaj' biçiminde satır yok.")
    return out


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

SEMA = {"type": "object", "properties": {"kayitlar": {"type": "array", "items": {"type": "object", "properties": {
    "tur": {"type": "string", "enum": ["Dönüş", "Söz", "Sorun"]}, "is_no": S, "tarih": S, "adet": S, "hatali_adet": S,
    "kalite_sonucu": {"type": "string", "enum": ["Kabul", "Ret", "Beklemede", ""]}, "hata_aciklamasi": S, "ozet": S, "kanit": S},
    "required": ["tur", "is_no", "tarih", "adet", "hatali_adet", "kalite_sonucu", "hata_aciklamasi", "ozet", "kanit"], "additionalProperties": False}}},
    "required": ["kayitlar"], "additionalProperties": False}


def denetle(m: dict, isler: dict, c: dict) -> list[str]:
    on, km, sorun = m["id"], katla(m["metin"]), []
    ns = sayilar(m["metin"])
    adaylar = aday_tarihler(m["metin"], m["gun"])
    gecen = [n for n in isler if re.search(rf"(?<![\w-]){re.escape(n)}(?![\w-])", m["metin"])]
    eksik = [n for n in gecen if n not in {k["is_no"] for k in c["kayitlar"]}]
    if eksik:
        sorun.append(f"{on}: mesajda geçen iş kayda alınmamış: {', '.join(eksik)}")
    for i, k in enumerate(c["kayitlar"], 1):
        ad = f"{on}: kayıt {i} ({k['tur']} {k['is_no']})"
        if k["is_no"] not in isler:
            sorun.append(f"{ad}: iş no listede yok")
        elif k["is_no"] not in gecen:
            sorun.append(f"{ad}: iş no mesajda geçmiyor")
        kk = katla(k["kanit"])
        if len(kk.split()) < 2 or kk not in km:
            sorun.append(f"{ad}: kanıt mesajda aynen geçmiyor")
        for alan_ad, v in (("adet", k["adet"]), ("hatalı adet", k["hatali_adet"])):
            if v.strip() and sayi(v) not in ns:
                sorun.append(f"{ad}: {alan_ad} '{v}' mesajda yazmıyor (hesap yapmayın; kalan adedi kod hesaplar)")
        if k["adet"].strip() and k["hatali_adet"].strip() and (sayi(k["hatali_adet"]) or 0) > (sayi(k["adet"]) or 0):
            sorun.append(f"{ad}: hatalı adet dönen adetten büyük")
        t = tarih_coz(k["tarih"])
        if not t:
            sorun.append(f"{ad}: tarih okunamadı ('{k['tarih']}')")
        elif k["tur"] == "Dönüş":
            if t > m["gun"]:
                sorun.append(f"{ad}: dönüş tarihi {t:%d.%m.%Y} mesaj tarihinden sonra; henüz teslim edilmemiş iş 'Söz' olmalı")
            elif t not in adaylar | {m["gun"]}:
                sorun.append(f"{ad}: dönüş tarihi {t:%d.%m.%Y} mesajdan çıkmıyor")
            if not TESLIM.search(kk):
                sorun.append(f"{ad}: kanıtta tamamlanmış teslim ifadesi yok (teslim edildi, gönderdik …); gelecek zamanlı ifade 'Söz' olmalı")
            if not k["adet"].strip():
                sorun.append(f"{ad}: dönüşte adet gerekli")
        elif k["tur"] == "Söz":
            if t not in adaylar:
                sorun.append(f"{ad}: söz tarihi {t:%d.%m.%Y} mesajdan çıkmıyor; mesaj tarihi {m['gun']:%d.%m.%Y}, mesajdan çıkan: "
                             + (", ".join(f"{x:%d.%m.%Y}" for x in sorted(adaylar)) or "yok"))
        elif t != m["gun"] and t not in adaylar:
            sorun.append(f"{ad}: sorun kaydının tarihi mesaj tarihi olmalı")
        if k["tur"] != "Dönüş" and (k["kalite_sonucu"] or k["hatali_adet"].strip()):
            sorun.append(f"{ad}: kalite sonucu / hatalı adet yalnız dönüşte yazılır")
    return sorun


def isler_metni(isler: dict) -> str:
    return "\n".join(f"{x['no']} · {x['atolye']} · {x['islem']} · çıkan {x['adet']:g} · termin {x['termin']:%d.%m.%Y}" if x["adet"] and x["termin"]
                     else f"{x['no']} · {x['atolye']} · {x['islem']}" for x in isler.values() if not x["kapandi"])


def uret(mesajlar: list[dict], isler: dict, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    il = isler_metni(isler)
    sonuc = dict(onceki or {})
    hedef = mesajlar if duzeltme is None else [m for m in mesajlar if any(s.startswith(m["id"] + ":") for s in duzeltme)]
    for m in hedef:
        mesaj = f"<isler>\n{il}\n</isler>\n<mesaj tarih=\"{m['zaman']}\" gonderen=\"{m['gonderen']}\">{m['metin']}</mesaj>"
        d = [s.split(": ", 1)[1] for s in (duzeltme or []) if s.startswith(m["id"] + ":")]
        if d:
            mesaj += ("\n<duzeltme>\nÖnceki cevap: " + json.dumps(sonuc.get(m["id"]), ensure_ascii=False) + "\nKodun bulduğu sorunlar:\n"
                      + "\n".join(f"- {s}" for s in d) + "\n</duzeltme>")
        sonuc[m["id"]] = llm.json_iste(sistem, mesaj, SEMA)
    return sonuc


def duzeltme_dongusu(uret_f, denetle_f, tur: int) -> tuple[dict, list[str], int]:
    taslak = uret_f(None, None)
    kalan = denetle_f(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret_f(kalan, taslak)
        kalan = denetle_f(taslak)
    return taslak, kalan, yapilan


# ----------------------------------------------------------------------------
# Hesap ve çıktı
# ----------------------------------------------------------------------------

def hesapla(mesajlar: list[dict], isler: dict, onceki: dict[str, Decimal], sonuc: dict) -> dict:
    """Dönüş toplamı, kalan adet, termin farkı ve aşım uyarıları (kod)."""
    donen = defaultdict(Decimal, onceki)
    notlar, sozler = defaultdict(list), []
    for m in mesajlar:
        for i, k in enumerate(sonuc[m["id"]]["kayitlar"]):
            x = isler.get(k["is_no"])
            if not x:
                continue
            if k["tur"] == "Dönüş" and (a := sayi(k["adet"])):
                donen[k["is_no"]] += a
                if x["adet"] and donen[k["is_no"]] > x["adet"]:
                    notlar[(m["id"], i)].append(f"Toplam dönüş {donen[k['is_no']]:g} çıkan adedi ({x['adet']:g}) aşıyor; mükerrer kayıt olabilir")
            if k["tur"] == "Söz":
                t = tarih_coz(k["tarih"])
                kalan = x["adet"] - donen[k["is_no"]] if x["adet"] is not None else None
                fark = (t - x["termin"]).days if t and x["termin"] else None
                if k["adet"].strip() and kalan is not None and sayi(k["adet"]) != kalan:
                    notlar[(m["id"], i)].append(f"Söz verilen adet {k['adet']}, kayıtlara göre kalan {kalan:g}")
                sozler.append({"is_no": k["is_no"], "atolye": x["atolye"], "termin": x["termin"], "soz": t, "fark": fark, "kalan": kalan,
                               "soz_adet": k["adet"], "mesaj": m["id"]})
    return {"donen": donen, "notlar": notlar, "sozler": sozler}


def tablo(ws, bas: list[str], satirlar: list[list], ai: set[int] = frozenset(), genis: dict | None = None) -> None:
    ws.append(bas)
    for h in ws[1]:
        h.fill, h.font, h.alignment = BASLIK_DOLGU, BASLIK_YAZI, UST
    for r in satirlar:
        ws.append(r)
        for j, v in enumerate(r, 1):
            c = ws.cell(ws.max_row, j)
            c.alignment = UST
            if isinstance(v, date):
                c.number_format = "DD.MM.YYYY"
            if j in ai and v not in (None, ""):
                c.fill = AI_DOLGU
    for j in range(1, len(bas) + 1):
        ws.column_dimensions[get_column_letter(j)].width = (genis or {}).get(j, 13)
    ws.freeze_panes = "A2"


def n(v: str) -> int | float | None:
    x = sayi(v)
    return None if x is None else (int(x) if x == x.to_integral_value() else float(x))


def rapor_yaz(cikti: Path, mesajlar: list[dict], isler: dict, sonuc: dict, h: dict, kalan: list[str], tur: int) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kayıtlar"
    rr = []
    for m in mesajlar:
        sor = "; ".join(x.split(": ", 1)[1] for x in kalan if x.startswith(m["id"] + ":"))
        ks = sonuc[m["id"]]["kayitlar"] or [None]
        for i, k in enumerate(ks):
            if k is None:
                rr.append([m["id"], m["zaman"], m["gonderen"], m["metin"], "", "", "Bilgi yok", None, None, None, "", "", "", "", sor or "Geçti"])
                continue
            rr.append([m["id"], m["zaman"], m["gonderen"], m["metin"], k["is_no"], isler.get(k["is_no"], {}).get("atolye", ""), k["tur"], tarih_coz(k["tarih"]),
                       n(k["adet"]), n(k["hatali_adet"]), k["kalite_sonucu"], k["hata_aciklamasi"], k["ozet"], " ".join(h["notlar"].get((m["id"], i), [])), sor or "Geçti"])
    tablo(ws, ["Mesaj", "Zaman", "Gönderen", "Mesaj Metni", "İş No", "Atölye", "Tür", "Tarih", "Adet", "Hatalı Adet", "Kalite", "Hata Açıklaması", "Özet",
               "Kod Notu", "Son Kontrol"], rr, ai={5, 7, 8, 9, 10, 11, 12, 13}, genis={4: 50, 12: 24, 13: 36, 14: 36, 15: 30})
    for i, r in enumerate(rr, 2):
        if r[13]:
            ws.cell(i, 14).fill = SARI
        ws.cell(i, 15).fill = YESIL if r[14] == "Geçti" else KIRMIZI
    g = wb.create_sheet("Sözler")
    tablo(g, ["İş No", "Atölye", "Termin", "Söz Verilen Tarih", "Termine Göre (gün)", "Kalan Adet (kod)", "Söz Verilen Adet", "Mesaj"],
          [[s["is_no"], s["atolye"], s["termin"], s["soz"], s["fark"], float(s["kalan"]) if s["kalan"] is not None else None, n(s["soz_adet"]), s["mesaj"]]
           for s in h["sozler"]], ai={4, 7})
    for i, s in enumerate(h["sozler"], 2):
        if s["fark"] is not None:
            g.cell(i, 5).fill = KIRMIZI if s["fark"] > 0 else YESIL
    k = wb.create_sheet("Kontrol")
    k.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    k.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    k.cell(2, 1).fill = KIRMIZI if kalan else YESIL
    k.append(["Mor hücreler modelin kayıtlarıdır; kalan adet, termin farkı ve aşım notları koddur. Dönüş kayıtları atölyenin beyanıdır: depo "
              "girişi ve kalite kontrol sonucuyla karşılaştırmadan hakedişe esas almayın. '_donusler.xlsx' dosyası 'Fason Atölye İş Takibi' kod "
              "bloğunun dönüş dosyasına eklenebilir."])
    k.column_dimensions["A"].width = 140
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    d = Workbook()
    tablo(d.active, ["İş No", "Dönüş Tarihi", "Dönen Adet", "Kalite Sonucu", "Hatalı Adet", "Hata Açıklaması"],
          [[k["is_no"], tarih_coz(k["tarih"]), n(k["adet"]), k["kalite_sonucu"] or "Beklemede", n(k["hatali_adet"]), k["hata_aciklamasi"]]
           for m in mesajlar for k in sonuc[m["id"]]["kayitlar"] if k["tur"] == "Dönüş" and k["is_no"] in isler])
    y = cikti.with_name(f"{cikti.stem}_donusler.xlsx")
    d.save(y)
    return y


def calistir(m_yol: Path, i_yol: Path, d_yol: Path | None, cikti: Path, gizle: list[str] | None = None, tur: int = 2, evet: bool = False) -> dict:
    isler, onceki, mesajlar = isleri_oku(i_yol), donusleri_oku(d_yol), mesajlari_oku(m_yol, gizle)
    print(f"[OK] {len(mesajlar)} mesaj · {len(isler)} fason iş · önceki dönüş kaydı olan iş {len(onceki)}")
    llm.onay_al(f"{len(mesajlar)} mesaj ve açık iş listesi (iş no, atölye, adet, termin) gönderilecek (telefon, e-posta, TCKN maskeli).", evet)
    sonuc, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(mesajlar, isler, d, o), lambda x: [s for m in mesajlar for s in denetle(m, isler, x[m["id"]])], tur)
    h = hesapla(mesajlar, isler, onceki, sonuc)
    dy = rapor_yaz(cikti, mesajlar, isler, sonuc, h, kalan, yapilan)
    return {"mesajlar": mesajlar, "sonuc": sonuc, "hesap": h, "kalan": kalan, "tur": yapilan, "donusler": dy}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Fason atölye mesajlarını dönüş, söz ve sorun kaydına dönüştürür.")
    p.add_argument("--mesajlar", type=Path, default=ORNEK / "mesajlar.txt", help="Mesajlar (.txt/.docx/.pdf): '[GG.AA.YYYY SS:DD] Gönderen: mesaj'")
    p.add_argument("--isler", type=Path, default=ORNEK / "fason_isler.csv", help="Fason iş listesi (Fason Atölye İş Takibi kod bloğunun dosyası)")
    p.add_argument("--donusler", type=Path, help="Önceki dönüşler (isteğe bağlı; kalan adet için)")
    p.add_argument("--gizle", nargs="*", default=[], metavar="AD", help="Mesajlarda geçen ve gönderilmeyecek kişi adları")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "fason_mesaj_kayitlari.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    d = a.donusler or (ORNEK / "donusler.csv" if a.mesajlar == ORNEK / "mesajlar.txt" else None)
    try:
        s = calistir(a.mesajlar, a.isler, d, a.cikti, a.gizle, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    for m in s["mesajlar"]:
        print(f"[OK] {m['id']} · " + (", ".join(f"{k['tur']} {k['is_no']} {k['tarih']}" + (f" {k['adet']} adet" if k["adet"] else "") for k in s["sonuc"][m["id"]]["kayitlar"]) or "kayıt yok"))
    for x in s["hesap"]["sozler"]:
        if x["fark"] is not None and x["fark"] > 0:
            print(f"[!] {x['is_no']} ({x['atolye']}): söz {x['soz']:%d.%m.%Y}, termin {x['termin']:%d.%m.%Y} → {x['fark']} gün gecikme")
    for x in s["kalan"]:
        print(f"[!] Kalan: {x}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[OK] Kod bloğu girdisi: {s['donusler'].name}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
