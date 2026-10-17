"""
Veli Bilgilendirme Mesajları — Gerisi Bende AI Agent
Kayıt Kabul ve Öğrenci İşleri › Öğrenci İşleri Sorumlusu

1. Kod mesaj listesinden mesaj türü kombinasyonlarını çıkarır: Devamsızlık, Ödev ve her duyuru (D1, D2 …) için bir şablon.
   Öğrenci ve veli bilgisi modele gönderilmez.
2. Model her kombinasyon için yer tutuculu kısa mesaj şablonu yazar: {VELI}, {OGRENCI}, {SINIF}, {OZURSUZ}, {OZURLU},
   {DERS}, {ODEV}, {OKUL}, {ILETISIM}. Duyuru şablonlarında duyurudaki tarih ve saatler aynen geçer.
3. Kod denetler:
   - Mesaj türünün zorunlu yer tutucuları var mı, bilinmeyen yer tutucu var mı
   - Duyurudaki her tarih ve saat şablonda aynen geçiyor mu, duyuruda olmayan sayı eklenmiş mi
   - Devamsızlık ve ödev şablonlarında sabit sayı yazılmış mı (sayılar yer tutucuyla gelir)
   - Yargılayıcı ya da tehdit içeren ifade ("tembel", "başarısız", "sınıfta kalacak", "disiplin") var mı
   - En uzun değerlerle doldurulan mesaj --azami karakteri (varsayılan 480) aşıyor mu
   Sorun varsa modele geri gönderilir (en fazla --tur kez). Kontrolden geçmeyen şablonla mesaj üretilmez.
4. Kod mesajları doldurur. Telefonu olmayan veliler "aranmalı" listesine alınır. SMS sistemine aktarılacak CSV üretilir;
   mesaj gönderilmez.

Kullanım:
    python agent.py                                        # örnek liste + duyurular
    python agent.py --girdi liste.xlsx --duyurular duyurular.txt --okul "Okul Adı" --iletisim "0212 000 00 00"
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
YER_TUTUCULAR = ["{VELI}", "{OGRENCI}", "{SINIF}", "{OZURSUZ}", "{OZURLU}", "{DERS}", "{ODEV}", "{OKUL}", "{ILETISIM}"]
ZORUNLU = {"Devamsızlık": ["{VELI}", "{OGRENCI}", "{OZURSUZ}", "{OZURLU}", "{OKUL}", "{ILETISIM}"],
           "Ödev": ["{VELI}", "{OGRENCI}", "{DERS}", "{ODEV}", "{OKUL}", "{ILETISIM}"],
           "Duyuru": ["{VELI}", "{OKUL}"]}
YASAK = re.compile(r"\b(tembel|basarisiz|sinifta kal|disiplin|ceza|ihmal|sorumsuz|ilgisiz)")
TARIH_SAAT = re.compile(r"\b\d{2}\.\d{2}\.\d{4}\b|\b\d{1,2}:\d{2}\b")
SEMA = {
    "type": "object",
    "properties": {"sablonlar": {"type": "array", "items": {"type": "object", "properties": {"kimlik": S, "metin": S},
                                                            "required": ["kimlik", "metin"], "additionalProperties": False}}},
    "required": ["sablonlar"],
    "additionalProperties": False,
}


def duyurulari_oku(yol: Path) -> dict[str, str]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: duyuru dosyası bulunamadı.")
    out = {}
    for s in belge.metin_oku(yol).splitlines():
        m = re.match(r"^\s*(D\d+)\s*[:.)-]\s*(.+)$", s)
        if m:
            out[m.group(1)] = m.group(2).strip()
    return out


def tur_norm(s: str) -> str:
    k = katla(s)
    return "Devamsızlık" if k.startswith("devam") else ("Ödev" if k.startswith("odev") else ("Duyuru" if k.startswith(("duyuru", "etkinlik")) else ""))


def listeyi_oku(yol: Path) -> list[dict]:
    r = kayitlar(yol, {"no": ("Öğrenci No", "Okul No", "No"), "ogrenci": ("Öğrenci Adı", "Öğrenci", "Ad Soyad"), "sinif": ("Sınıf", "Şube"),
                       "veli": ("Veli Adı", "Veli"), "tel": ("Veli Telefon", "Telefon", "GSM"), "tur": ("Mesaj Türü", "Tür"), "duyuru": ("Duyuru", "Duyuru No"),
                       "ozursuz": ("Özürsüz Gün", "Özürsüz"), "ozurlu": ("Özürlü Gün", "Özürlü"), "ders": ("Ders",), "odev": ("Eksik Ödev", "Ödev Sayısı")},
                 ("ogrenci", "tur"))
    out = []
    for x in r:
        if not x["ogrenci"]:
            continue
        out.append({k: (str(v).strip() if not isinstance(v, (int, float)) else sayi_yaz(v)) for k, v in x.items()} | {"tur": tur_norm(x["tur"])})
    return out


def kombinasyonlar(liste: list[dict], duyurular: dict[str, str]) -> list[dict]:
    out, gorulen = [], set()
    for x in liste:
        anahtar = (x["tur"], x["duyuru"] if x["tur"] == "Duyuru" else "")
        if not x["tur"] or anahtar in gorulen:
            continue
        gorulen.add(anahtar)
        out.append({"kimlik": f"S{len(out) + 1}", "tur": x["tur"], "duyuru": anahtar[1], "duyuru_metni": duyurular.get(anahtar[1], "")})
    return out


def doldur(sablon: str, x: dict, okul: str, iletisim: str) -> str:
    deger = {"{VELI}": x["veli"] or "Sayın Velimiz", "{OGRENCI}": x["ogrenci"], "{SINIF}": x["sinif"], "{OZURSUZ}": x["ozursuz"] or "0", "{OZURLU}": x["ozurlu"] or "0",
             "{DERS}": x["ders"], "{ODEV}": x["odev"] or "0", "{OKUL}": okul, "{ILETISIM}": iletisim}
    for k, v in deger.items():
        sablon = sablon.replace(k, v)
    return re.sub(r"\s{2,}", " ", sablon).strip()


def denetle(t: dict, komb: list[dict], liste: list[dict], okul: str, iletisim: str, azami: int) -> list[str]:
    sorun, sab = [], {s["kimlik"]: s["metin"] for s in t["sablonlar"]}
    for k in komb:
        ad = f"{k['kimlik']} ({k['tur']}{' ' + k['duyuru'] if k['duyuru'] else ''})"
        m = sab.get(k["kimlik"])
        if not m:
            sorun.append(f"{ad}: şablon yazılmadı")
            continue
        eksik = [y for y in ZORUNLU[k["tur"]] if y not in m]
        if eksik:
            sorun.append(f"{ad}: zorunlu yer tutucu eksik: {', '.join(eksik)}")
        bilinmeyen = [y for y in re.findall(r"\{[^}]*\}", m) if y not in YER_TUTUCULAR]
        if bilinmeyen:
            sorun.append(f"{ad}: bilinmeyen yer tutucu: {', '.join(bilinmeyen)}")
        ekli = re.findall(r"\{[A-Z_]+\}['’]?[a-zçğıöşü]+", m)
        if ekli:
            sorun.append(f"{ad}: yer tutucuya ek getirilmiş ({', '.join(ekli)}); ad ve ders adına uymayabilir, cümleyi eksiz kurun")
        sade = re.sub(r"\{[^}]*\}", " ", m)
        if k["tur"] == "Duyuru":
            for z in TARIH_SAAT.findall(k["duyuru_metni"]):
                if z not in sade:
                    sorun.append(f"{ad}: duyurudaki '{z}' şablonda aynen geçmiyor")
            fazla = sorted(sayilar(sade) - sayilar(k["duyuru_metni"]))
            if fazla:
                sorun.append(f"{ad}: duyuruda olmayan sayı / tarih: {', '.join(fazla)}")
        elif sayilar(sade):
            sorun.append(f"{ad}: sabit sayı yazılmış ({', '.join(sorted(sayilar(sade)))}); sayılar yer tutucuyla gelmeli")
        if YASAK.search(katla(sade)):
            sorun.append(f"{ad}: yargılayıcı ya da tehdit içeren ifade ('{YASAK.search(katla(sade)).group(0)}')")
        ilgili = [x for x in liste if x["tur"] == k["tur"] and (k["tur"] != "Duyuru" or x["duyuru"] == k["duyuru"])]
        if ilgili:
            en_uzun = max(len(doldur(m, x, okul, iletisim)) for x in ilgili)
            if en_uzun > azami:
                sorun.append(f"{ad}: doldurulmuş mesaj {en_uzun} karakter, sınır {azami}")
    return sorun


def uret(komb: list[dict], azami: int, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    satir = "\n".join(f"{k['kimlik']} | tür: {k['tur']} | zorunlu: {' '.join(ZORUNLU[k['tur']])}" + (f" | duyuru: {k['duyuru_metni']}" if k["duyuru"] else "") for k in komb)
    mesaj = f"<kombinasyonlar>\n{satir}\n</kombinasyonlar>\n<azami_karakter>{azami}</azami_karakter>" + duzeltme_blogu(duzeltme, onceki)
    return llm.json_iste(sistem, mesaj, SEMA)


def rapor_yaz(cikti: Path, komb, t, mesajlar, atlanan, kalan, tur) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Şablonlar"
    sab = {s["kimlik"]: s["metin"] for s in t["sablonlar"]}
    baslik(ws, ["Kimlik", "Tür", "Duyuru", "Şablon", "Kontrol", "Onay"], (7, 12, 50, 80, 30, 8))
    for k in komb:
        sor = [s for s in kalan if s.startswith(k["kimlik"] + " ")]
        ws.append([k["kimlik"], k["tur"], k["duyuru_metni"], sab.get(k["kimlik"], ""), "; ".join(sor) or "Geçti", ""])
        boya(ws, [4], AI_DOLGU)
        boya(ws, [5], KIRMIZI if sor else YESIL)
        boya(ws, [6], SARI)
    son_not(ws, "Önemli: Şablonları okul idaresi onaylamalıdır. Mesajlar gönderilmedi; CSV dosyası SMS sistemine aktarılmak içindir.", kalan, tur)
    hizala(ws)
    m = wb.create_sheet("Mesajlar")
    baslik(m, ["Öğrenci No", "Öğrenci", "Sınıf", "Veli", "Telefon", "Tür", "Mesaj", "Karakter"], (10, 18, 7, 16, 16, 12, 90, 9))
    for x, metin in mesajlar:
        m.append([x["no"], x["ogrenci"], x["sinif"], x["veli"], x["tel"], x["tur"], metin, len(metin)])
    hizala(m)
    a = wb.create_sheet("Gönderilemeyen")
    baslik(a, ["Öğrenci No", "Öğrenci", "Veli", "Tür", "Neden"], (10, 18, 16, 12, 60))
    for x, neden in atlanan:
        a.append([x["no"], x["ogrenci"], x["veli"], x["tur"], neden])
    hizala(a)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    sms = cikti.with_name(cikti.stem + "_sms.csv")
    with sms.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Telefon", "Mesaj"])
        for x, metin in mesajlar:
            w.writerow([x["tel"], metin])
    return sms


def calistir(girdi: Path, duyuru_yolu: Path, cikti: Path, okul: str, iletisim: str, azami: int = 480, tur: int = 2, evet: bool = False) -> dict:
    liste = listeyi_oku(girdi)
    duyurular = duyurulari_oku(duyuru_yolu)
    eksik = sorted({x["duyuru"] for x in liste if x["tur"] == "Duyuru" and x["duyuru"] not in duyurular})
    if eksik:
        raise llm.LLMHatasi(f"Duyuru dosyasında olmayan duyuru numarası: {', '.join(eksik)}")
    komb = kombinasyonlar(liste, duyurular)
    if not komb:
        raise llm.LLMHatasi(f"{girdi.name}: mesaj türü bulunamadı (Devamsızlık / Ödev / Duyuru).")
    print(f"[OK] {len(liste)} satır · {len(komb)} şablon kombinasyonu")
    llm.onay_al(f"{len(komb)} mesaj türü ve duyuru metni gönderilecek (öğrenci ve veli bilgisi gönderilmez).", evet)
    t, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(komb, azami, d, o), lambda x: denetle(x, komb, liste, okul, iletisim, azami), tur)
    sab = {s["kimlik"]: s["metin"] for s in t["sablonlar"]}
    gecmeyen = {k["kimlik"] for k in komb if any(s.startswith(k["kimlik"] + " ") for s in kalan)}
    kimlik = {(k["tur"], k["duyuru"]): k["kimlik"] for k in komb}
    mesajlar, atlanan = [], []
    for x in liste:
        if not x["tur"]:
            atlanan.append((x, "Mesaj türü tanınmadı"))
            continue
        kid = kimlik[(x["tur"], x["duyuru"] if x["tur"] == "Duyuru" else "")]
        if kid in gecmeyen or kid not in sab:
            atlanan.append((x, f"Şablon {kid} kontrolden geçmedi"))
        elif not x["tel"]:
            atlanan.append((x, "Veli telefonu yok; aranmalı"))
        else:
            mesajlar.append((x, doldur(sab[kid], x, okul, iletisim)))
    sms = rapor_yaz(cikti, komb, t, mesajlar, atlanan, kalan, yapilan)
    return {"liste": liste, "komb": komb, "taslak": t, "mesajlar": mesajlar, "atlanan": atlanan, "kalan": kalan, "tur": yapilan, "sms": sms}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    p = argparse.ArgumentParser(description="Duyuru, devamsızlık ve ödev bilgilerinden kişiselleştirilmiş veli mesajları hazırlar (göndermez).")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "mesaj_listesi.csv",
                   help="Mesaj listesi (.xlsx/.csv): Öğrenci Adı, Mesaj Türü; Öğrenci No, Sınıf, Veli Adı, Veli Telefon, Duyuru, Özürsüz Gün, Özürlü Gün, Ders, Eksik Ödev")
    p.add_argument("--duyurular", type=Path, default=BURASI / "ornek_veri" / "duyurular.txt", help="Duyurular (.txt/.docx): 'D1: …' satırları")
    p.add_argument("--okul", default="Örnek Anadolu Lisesi", help="Okul adı")
    p.add_argument("--iletisim", default="0212 000 00 00", help="Velilerin arayabileceği okul telefonu")
    p.add_argument("--azami", type=int, default=480, help="Doldurulmuş mesajın en fazla karakteri (varsayılan 480)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "veli_mesajlari.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.duyurular, a.cikti, a.okul, a.iletisim, a.azami, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(s['mesajlar'])} mesaj hazır (gönderilmedi) · gönderilemeyen {len(s['atlanan'])}")
    for x, neden in s["atlanan"]:
        print(f"[!] {x['ogrenci']} ({x['tur']}): {neden}")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()} (+ {s['sms'].name})")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
