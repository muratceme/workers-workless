"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

LISTE = KLASOR / "ornek_veri" / "mesaj_listesi.csv"
DUYURU = KLASOR / "ornek_veri" / "duyurular.txt"
OKUL, TEL = "Örnek Anadolu Lisesi", "0212 000 00 00"
MESAJLAR = []
DOGRU = {
    "S1": "Sayın {VELI}, öğrencimiz {OGRENCI} ({SINIF}) bugüne kadar {OZURSUZ} gün özürsüz, {OZURLU} gün özürlü devamsızlık yapmıştır. "
          "Sorularınız için {ILETISIM} numaralı telefondan bize ulaşabilirsiniz. {OKUL}",
    "S2": "Sayın {VELI}, öğrencimiz {OGRENCI} için {DERS} dersinde eksik ödev sayısı: {ODEV}. Desteğiniz için teşekkür ederiz. Bilgi: {ILETISIM}. {OKUL}",
    "S3": "Sayın {VELI}, veli toplantımız 24.10.2026 Cumartesi saat 10:00'da okul konferans salonunda yapılacaktır. Katılımınızı rica ederiz. {OKUL}",
    "S4": "Sayın {VELI}, Cumhuriyet Bayramı töreni 29.10.2026 saat 09:30'da okul bahçesinde yapılacaktır; öğrencilerimizin 09:00'da okulda olması "
          "gerekmektedir. {OKUL}",
}


def sablonlar(dogru=True) -> dict:
    d = dict(DOGRU)
    if not dogru:
        d["S1"] = d["S1"].replace("yapmıştır.", "yapmıştır. 10 günü aşarsa sınıfta kalacaktır.")
        d["S2"] = d["S2"].replace("öğrencimiz {OGRENCI} için", "öğrencimiz {OGRENCI}'in").replace(" Bilgi: {ILETISIM}.", "")
        d["S3"] = d["S3"].replace("10:00", "10:30")
    return {"sablonlar": [{"kimlik": k, "metin": v} for k, v in d.items()]}


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return sablonlar("<duzeltme>" in kullanici)


class KodTesti(unittest.TestCase):
    def test_kombinasyon(self):
        l, d = agent.listeyi_oku(LISTE), agent.duyurulari_oku(DUYURU)
        k = agent.kombinasyonlar(l, d)
        self.assertEqual([(x["kimlik"], x["tur"], x["duyuru"]) for x in k], [("S1", "Devamsızlık", ""), ("S2", "Ödev", ""), ("S3", "Duyuru", "D1"), ("S4", "Duyuru", "D2")])
        self.assertEqual(agent.doldur(DOGRU["S1"], l[0], OKUL, TEL)[:62], "Sayın Selma Aydın, öğrencimiz Deniz Aydın (9-A) bugüne kadar 4")

    def test_denetim(self):
        l, d = agent.listeyi_oku(LISTE), agent.duyurulari_oku(DUYURU)
        k = agent.kombinasyonlar(l, d)
        self.assertEqual(agent.denetle(sablonlar(), k, l, OKUL, TEL, 480), [])
        s = agent.denetle(sablonlar(False), k, l, OKUL, TEL, 480)
        self.assertEqual(len(s), 6)
        self.assertTrue(any("S1" in x and "sabit sayı" in x for x in s) and any("S1" in x and "sinifta kal" in x for x in s)
                        and any("S2" in x and "{ILETISIM}" in x for x in s) and any("S2" in x and "ek getirilmiş" in x for x in s)
                        and any("'10:00'" in x for x in s) and any("10:30" in x for x in s))
        self.assertTrue(any("sınır 100" in x for x in agent.denetle(sablonlar(), k, l, OKUL, TEL, 100)))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "v.xlsx"
            s = agent.calistir(LISTE, DUYURU, cikti, OKUL, TEL, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            for gizli in ("Deniz", "Selma", "0532", "1021", "9-A"):
                self.assertNotIn(gizli, MESAJLAR[0])
            self.assertEqual((len(s["mesajlar"]), [(x["ogrenci"], n) for x, n in s["atlanan"]]), (7, [("Ece Şahin", "Veli telefonu yok; aranmalı")]))
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Şablonlar", "Mesajlar", "Gönderilemeyen"])
            self.assertEqual(len(s["sms"].read_text(encoding="utf-8-sig").splitlines()), 8)

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(LISTE, DUYURU, Path(tmp) / "v.xlsx", OKUL, TEL, tur=0, evet=True)
            self.assertEqual((len(s["kalan"]), len(s["mesajlar"]), len(s["atlanan"])), (6, 1, 7))
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--duyurular", str(DUYURU.with_name("yok.txt")), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
