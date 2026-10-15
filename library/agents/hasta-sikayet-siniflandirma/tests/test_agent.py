"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import re
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

BAS = KLASOR / "ornek_veri" / "basvurular.csv"
BIRIM = KLASOR / "ornek_veri" / "birimler.txt"
BUGUN = date(2026, 10, 15)
MESAJLAR = []
DOGRU = {  # id: (birim, konu, tür, önem, alıntı)
    "HB-300": ("Laboratuvar", "Bekleme süresi / randevu", "Şikâyet", "Orta", "3 gündür sisteme girilmedi"),
    "HB-301": ("Kardiyoloji Polikliniği", "Bekleme süresi / randevu", "Şikâyet", "Orta", "2 saat sonra muayene oldum"),
    "HB-302": ("Acil Servis", "Şiddet / güvenlik", "Şikâyet", "Yüksek", "personel tehdit edildi"),
    "HB-303": ("Dahiliye", "İletişim / davranış", "Şikâyet", "Orta", "soru sorunca kaba davrandı"),
    "HB-304": ("Genel Cerrahi Servisi", "Hasta güvenliği olayı", "Şikâyet", "Yüksek", "yataktan düştü"),
    "HB-305": ("Hasta Kabul ve Faturalama", "Ücret / faturalama", "Şikâyet", "Orta", "hiç almadığım bir ultrason işlemi yazılmış"),
    "HB-306": ("Hasta Hizmetleri / Danışma", "Fiziki ortam / erişilebilirlik", "Öneri", "Düşük", "su sebili konulursa"),
    "HB-307": ("Kadın Doğum Servisi", "İletişim / davranış", "Teşekkür", "Düşük", "çok teşekkür ederim"),
    "HB-308": ("Hasta Hizmetleri / Danışma", "Mahremiyet / gizlilik", "Şikâyet", "Yüksek", "yüksek sesle okundu"),
    "HB-309": ("Otelcilik (temizlik, yemek)", "Temizlik / otelcilik / yemek", "Şikâyet", "Orta", "yemeği soğuk geliyor"),
}


def sonuc(dogru=True) -> dict:
    d = {i: {"id": i, "birim": b, "konu": k, "tur": t, "onem": o, "ozet": "özet", "alinti": a, "inceleme_sorusu": "?"} for i, (b, k, t, o, a) in DOGRU.items()}
    if not dogru:
        d["HB-304"]["onem"] = "Orta"
        d["HB-308"]["birim"] = "Poliklinik Bekleme"
        d["HB-306"]["alinti"] = "su sebili konmalı"
    return d


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    d = sonuc("<duzeltme>" in kullanici)
    return {"basvurular": [d[i] for i in re.findall(r'<basvuru id="([^"]+)"', kullanici)]}


class KodTesti(unittest.TestCase):
    def test_kural_ve_ad(self):
        b = agent.basvurulari_oku(BAS)
        self.assertEqual({x["id"]: agent.kural(x["metin"]) for x in b if agent.kural(x["metin"])},
                         {"HB-302": ["şiddet / tehdit"], "HB-304": ["hasta güvenliği olayı"], "HB-308": ["mahremiyet"]})
        self.assertEqual(agent.gonderilecek(b[3]), "Murat Koç adlı doktor bana tahlil sonuçlarımı açıklamadı, soru sorunca kaba davrandı. [HASTA]")
        self.assertEqual(agent.adi_cikar("Ben Ayşe Kaya, ayşe hanım diye seslendiler", "Ayşe Kaya"), "Ben [HASTA] , [HASTA] hanım diye seslendiler")

    def test_sure(self):
        b = {x["id"]: x for x in agent.basvurulari_oku(BAS)}
        s = {k: agent.sure_durumu(v, 30, BUGUN, 7) for k, v in b.items()}
        self.assertEqual([s[k][1:] for k in ("HB-300", "HB-301", "HB-302", "HB-303")],
                         [("GECİKTİ (14 gün)", -14), ("Zamanında cevaplandı", None), ("Süre doluyor", 7), ("Süresinde", 13)])
        self.assertEqual(agent.sure_durumu(b["HB-301"], 10, BUGUN, 7)[1], "Geç cevaplandı (5 gün)")

    def test_denetim(self):
        b = agent.basvurulari_oku(BAS)
        birimler = BIRIM.read_text(encoding="utf-8").split("\n")
        self.assertEqual(agent.denetle(b, sonuc(), birimler), [])
        s = agent.denetle(b, sonuc(False), birimler)
        self.assertEqual(len(s), 3)
        self.assertTrue(any("HB-304" in x and "Yüksek" in x for x in s) and any("Poliklinik Bekleme" in x for x in s) and any("HB-306" in x for x in s))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "h.xlsx"
            s = agent.calistir(BAS, BIRIM, cikti, bugun=BUGUN, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            for ad in ("Selma", "Arslan", "Elif Şahin", "Hakan", "Zeynep"):
                self.assertNotIn(ad, MESAJLAR[0])
            self.assertEqual(re.findall(r'<basvuru id="([^"]+)"', MESAJLAR[1]), ["HB-304", "HB-306", "HB-308"])
            ws = load_workbook(cikti)["Başvurular"]
            self.assertEqual([ws.cell(r, 1).value for r in range(2, 5)], ["HB-302", "HB-304", "HB-308"])
            self.assertEqual(ws.cell(5, 16).value, "GECİKTİ (14 gün)")

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(BAS, BIRIM, Path(tmp) / "h.xlsx", bugun=BUGUN, tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 3)
            self.assertEqual(agent.main(["--evet", "--bugun", "15.10.2026", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--bugun", "x"]), 2)
            self.assertEqual(agent.main(["--girdi", str(BAS.with_name("yok.csv")), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
