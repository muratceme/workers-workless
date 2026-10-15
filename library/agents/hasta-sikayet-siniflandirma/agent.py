"""
Hasta Şikâyet Sınıflandırma — Workers / Workless AI Agent
Hasta Hizmetleri › Hasta İlişkileri ve Hasta Hakları Yöneticisi

1. Kod başvuruları okur; satırdaki hasta adını metinden çıkarır ([HASTA]), telefon / e-posta / TCKN'yi maskeler. Önem
   kurallarını uygular: şiddet / tehdit / hakaret, hasta güvenliği olayı (düşme, yanlış ilaç, yanlış taraf, enfeksiyon,
   yanık) ya da mahremiyet ihlali belirtisi olan başvuru "Yüksek" önemdedir.
2. Model her başvuruyu birim (verilen listeden), konu, tür (şikâyet / öneri / teşekkür / bilgi talebi), önem ve kısa
   özete ayırır; dayandığı ifadeyi alıntılar ve ilgili birime iletilecek inceleme sorusunu yazar.
3. Kod denetler: birim listede mi, alıntı metinde mi, kural gereği Yüksek olması gereken başvuru Yüksek mi. Sorun varsa
   modele geri gönderilir (en fazla --tur kez).
4. Kod süreleri izler: son cevap tarihi = başvuru tarihi + --sure gün (varsayılan 30; kurum prosedürünüze göre
   ayarlayın). Açık başvurularda kalan gün, --uyari gün (varsayılan 7) içinde kalanlar ve gecikenler; kapalılarda
   zamanında cevap verilip verilmediği.

Kullanım:
    python agent.py                                               # örnek başvurular + birim listesi
    python agent.py --girdi basvurular.xlsx --birimler birimler.txt --sure 30 --bugun 15.10.2026
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
PAKET = 25
KONULAR = ["Tıbbi bakım / tedavi", "İletişim / davranış", "Bekleme süresi / randevu", "Bilgilendirme", "Mahremiyet / gizlilik", "Ücret / faturalama",
           "Temizlik / otelcilik / yemek", "Şiddet / güvenlik", "Hasta güvenliği olayı", "Kayıt / evrak", "Fiziki ortam / erişilebilirlik", "Diğer"]
TURLER = ["Şikâyet", "Öneri", "Teşekkür", "Bilgi talebi"]
ONEM = ["Yüksek", "Orta", "Düşük"]
KURALLAR = {
    "şiddet / tehdit": re.compile(r"\b(siddet|darp|tehdit|hakaret|vurdu|saldir|taciz|kufur)"),
    "hasta güvenliği olayı": re.compile(r"\b(dustu|dusme|yanlis ilac|yanlis taraf|yanlis hasta|enfeksiyon|yanik|ilac hatasi)"),
    "mahremiyet": re.compile(r"(mahremiyet|gizlilik|herkes duydu|yuksek sesle okundu|herkesin onunde|ifsa|izinsiz paylas)"),
}
SEMA = {
    "type": "object",
    "properties": {"basvurular": {"type": "array", "items": {"type": "object", "properties": {
        "id": S, "birim": S, "konu": {"type": "string", "enum": KONULAR}, "tur": {"type": "string", "enum": TURLER}, "onem": {"type": "string", "enum": ONEM},
        "ozet": S, "alinti": S, "inceleme_sorusu": S},
        "required": ["id", "birim", "konu", "tur", "onem", "ozet", "alinti", "inceleme_sorusu"], "additionalProperties": False}}},
    "required": ["basvurular"],
    "additionalProperties": False,
}
YAPILACAK = {"şiddet / tehdit": "Çalışan güvenliği (beyaz kod) prosedürü ve güvenlik kayıtları kontrol edilmeli",
             "hasta güvenliği olayı": "Hasta güvenliği olay bildirimi yapılıp yapılmadığı kontrol edilmeli",
             "mahremiyet": "Mahremiyet ihlali yönünden ilgili birim yöneticisine ve KVKK sorumlusuna iletilmeli"}


def kural(metin: str) -> list[str]:
    k = katla(metin)
    return [ad for ad, d in KURALLAR.items() if d.search(k)]


def adi_cikar(metin: str, ad: str) -> str:
    out = metin
    for parca in sorted({ad.strip(), *[p for p in ad.split() if len(p) >= 3]}, key=len, reverse=True):
        if parca:
            out = re.sub(re.escape(parca), "[HASTA]", out, flags=re.I)
    return re.sub(r"(\[HASTA\]\s*)+", "[HASTA] ", out).strip()


def basvurulari_oku(yol: Path) -> list[dict]:
    r = kayitlar(yol, {"id": ("Başvuru No", "Kayıt No", "No"), "tarih": ("Başvuru Tarihi", "Tarih"), "kanal": ("Kanal", "Başvuru Şekli"),
                       "hasta": ("Hasta Adı", "Başvuran", "Ad Soyad"), "metin": ("Metin", "Başvuru", "Açıklama", "Şikâyet"),
                       "durum": ("Durum",), "cevap": ("Cevap Tarihi", "Kapanış Tarihi", "Yanıt Tarihi")}, ("metin", "tarih"))
    return [{"id": str(x["id"] or f"HB{i}"), "tarih": tarih_coz(x["tarih"]), "kanal": str(x["kanal"]), "hasta": str(x["hasta"]), "metin": str(x["metin"]),
             "durum": str(x["durum"] or "Açık"), "cevap": tarih_coz(x["cevap"])} for i, x in enumerate(r, 1) if str(x["metin"]).strip()]


def gonderilecek(b: dict) -> str:
    return llm.maskele(adi_cikar(b["metin"], b["hasta"]) if b["hasta"] else b["metin"])


def denetle(bas: list[dict], sonuc: dict, birimler: list[str]) -> list[str]:
    sorun = []
    for b in bas:
        s = sonuc.get(b["id"])
        if not s:
            sorun.append(f"{b['id']}: sınıflandırılmadı")
            continue
        if s["birim"] not in birimler:
            sorun.append(f"{b['id']}: birim '{s['birim']}' listede yok (listeden seçin)")
        if not metinde_var(s["alinti"], gonderilecek(b)):
            sorun.append(f"{b['id']}: alıntı metinde yok: '{s['alinti'][:50]}'")
        k = kural(b["metin"])
        if k and s["onem"] != "Yüksek":
            sorun.append(f"{b['id']}: metinde {', '.join(k)} belirtisi var; önem 'Yüksek' olmalı ({s['onem']} verilmiş)")
    return sorun


def uret(bas: list[dict], birimler: list[str], duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    sonuc, hedef = {}, bas
    if duzeltme is not None:  # düzeltmede yalnız sorunlu başvurular yeniden sorulur
        sorunlu = {s.split(":")[0] for s in duzeltme}
        sonuc.update({a: b for a, b in (onceki or {}).items() if a not in sorunlu})
        hedef = [b for b in bas if b["id"] in sorunlu]
    for j in range(0, len(hedef), PAKET):
        parca = hedef[j:j + PAKET]
        ids = {b["id"] for b in parca}
        satir = "\n".join(f'<basvuru id="{b["id"]}" kanal="{b["kanal"]}" kod_bulgusu="{", ".join(kural(b["metin"])) or "-"}">{gonderilecek(b)}</basvuru>' for b in parca)
        d = [s for s in (duzeltme or []) if s.split(":")[0] in ids]
        mesaj = (f"<birimler>{'; '.join(birimler)}</birimler>\n<konular>{'; '.join(KONULAR)}</konular>\n<basvurular>\n{satir}\n</basvurular>"
                 + duzeltme_blogu(d, [v for a, v in (onceki or {}).items() if a in ids]))
        for s in llm.json_iste(sistem, mesaj, SEMA).get("basvurular", []):
            if s.get("id") in ids:
                sonuc[s["id"]] = s
    return sonuc


def sure_durumu(b: dict, sure: int, bugun: date, uyari: int) -> tuple[date | None, str, int | None]:
    """(son cevap tarihi, durum metni, kalan gün)"""
    if not b["tarih"]:
        return None, "Başvuru tarihi yok", None
    son = b["tarih"] + timedelta(days=sure)
    if katla(b["durum"]) in ("kapali", "kapandi", "cevaplandi", "sonuclandi") or b["cevap"]:
        if b["cevap"]:
            return son, ("Zamanında cevaplandı" if b["cevap"] <= son else f"Geç cevaplandı ({(b['cevap'] - son).days} gün)"), None
        return son, "Kapalı (cevap tarihi yok)", None
    kalan = (son - bugun).days
    if kalan < 0:
        return son, f"GECİKTİ ({-kalan} gün)", kalan
    return son, ("Süre doluyor" if kalan <= uyari else "Süresinde"), kalan


def rapor_yaz(cikti: Path, bas, sonuc, sureler, birimler, kalan, tur, sure) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Başvurular"
    baslik(ws, ["Başvuru No", "Tarih", "Kanal", "Tür", "Önem", "Birim", "Konu", "Özet", "Alıntı", "Kod Bulgusu", "Yapılacak", "İnceleme Sorusu (birime)",
                "Durum", f"Son Cevap ({sure} gün)", "Kalan Gün", "Süre Durumu", "Sorumlu"], (10, 11, 13, 10, 8, 20, 22, 40, 30, 20, 40, 40, 8, 12, 8, 18, 12))
    sira = sorted(bas, key=lambda b: (ONEM.index(sonuc[b["id"]]["onem"]) if b["id"] in sonuc else 9, sureler[b["id"]][2] if sureler[b["id"]][2] is not None else 999))
    for b in sira:
        s = sonuc.get(b["id"])
        son, durum, kal = sureler[b["id"]]
        k = kural(b["metin"])
        if not s:
            ws.append([b["id"], b["tarih"], b["kanal"], "", "", "", "", "AI yanıt vermedi"])
            continue
        ws.append([b["id"], b["tarih"], b["kanal"], s["tur"], s["onem"], s["birim"], s["konu"], s["ozet"], s["alinti"], ", ".join(k),
                   "; ".join(YAPILACAK[x] for x in k), s["inceleme_sorusu"], b["durum"], son, kal, durum, ""])
        ws.cell(ws.max_row, 2).number_format = ws.cell(ws.max_row, 14).number_format = "DD.MM.YYYY"
        boya(ws, [4, 6, 7, 8, 9, 12], AI_DOLGU)
        if s["onem"] == "Yüksek":
            boya(ws, [5], KIRMIZI)
        boya(ws, [16], KIRMIZI if durum.startswith(("GECİKTİ", "Geç")) else (SARI if durum == "Süre doluyor" else YESIL))
        boya(ws, [17], SARI)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:Q{ws.max_row}"
    son_not(ws, "Önemli: Sınıflandırma yapay zekâ önerisidir. Hasta hakları birimi her başvuruyu okuyarak karar verir; cevap metni ayrıca hazırlanır.", kalan, tur)
    hizala(ws)
    m = wb.create_sheet("Birim × Konu")
    sik = [b for b in bas if b["id"] in sonuc and sonuc[b["id"]]["tur"] == "Şikâyet"]
    konular = [k for k in KONULAR if any(sonuc[b["id"]]["konu"] == k for b in sik)]
    baslik(m, ["Birim", *konular, "Toplam Şikâyet"], (26, *[14] * len(konular), 14))
    say = Counter((sonuc[b["id"]]["birim"], sonuc[b["id"]]["konu"]) for b in sik)
    for br in birimler:
        sat = [say[(br, k)] for k in konular]
        if sum(sat):
            m.append([br, *[v or "" for v in sat], sum(sat)])
    m.append([])
    t = Counter(sonuc[b["id"]]["tur"] for b in bas if b["id"] in sonuc)
    m.append(["Başvuru türü"] + [f"{k}: {t[k]}" for k in TURLER])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(girdi: Path, birim_yolu: Path, cikti: Path, sure: int = 30, uyari: int = 7, bugun: date | None = None, tur: int = 2, evet: bool = False) -> dict:
    bas = basvurulari_oku(girdi)
    if not bas:
        raise llm.LLMHatasi(f"{girdi.name}: başvuru yok.")
    if not birim_yolu.exists():
        raise llm.LLMHatasi(f"{birim_yolu}: birim listesi bulunamadı.")
    birimler = [s.strip() for s in birim_yolu.read_text(encoding="utf-8-sig").splitlines() if s.strip()]
    if "Belirsiz" not in birimler:
        birimler.append("Belirsiz")
    bugun = bugun or date.today()
    print(f"[OK] {len(bas)} başvuru · kural bulgusu {sum(1 for b in bas if kural(b['metin']))} · {len(birimler)} birim")
    llm.onay_al(f"{len(bas)} başvuru metni gönderilecek. Hasta adı metinden çıkarılır; telefon / e-posta / TCKN maskelenir. "
                "Metindeki diğer kişi adları ve sağlık bilgileri gönderilir.", evet)
    sonuc, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(bas, birimler, d, o), lambda x: denetle(bas, x, birimler), tur)
    sureler = {b["id"]: sure_durumu(b, sure, bugun, uyari) for b in bas}
    rapor_yaz(cikti, bas, sonuc, sureler, birimler, kalan, yapilan, sure)
    return {"basvurular": bas, "sonuc": sonuc, "sureler": sureler, "kalan": kalan, "tur": yapilan}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    p = argparse.ArgumentParser(description="Hasta şikâyet ve önerilerini birim, konu ve önem derecesine göre sınıflandırıp cevap süresiyle takip eder.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "basvurular.csv",
                   help="Başvurular (.xlsx/.csv): Başvuru Tarihi, Metin; Başvuru No, Kanal, Hasta Adı, Durum, Cevap Tarihi isteğe bağlı")
    p.add_argument("--birimler", type=Path, default=BURASI / "ornek_veri" / "birimler.txt", help="Birim listesi (.txt, her satırda bir birim)")
    p.add_argument("--sure", type=int, default=30, help="Cevap süresi, gün (varsayılan 30; kurum prosedürünüze göre ayarlayın)")
    p.add_argument("--uyari", type=int, default=7, help="Sürenin dolmasına bu kadar gün kalanları işaretle (varsayılan 7)")
    p.add_argument("--bugun", help="Hesap tarihi GG.AA.YYYY (varsayılan bugün)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "hasta_basvurulari.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    bugun = tarih_coz(a.bugun) if a.bugun else None
    if a.bugun and not bugun:
        print(f"[X] Tarih anlaşılamadı: {a.bugun}")
        return 2
    try:
        s = calistir(a.girdi, a.birimler, a.cikti, a.sure, a.uyari, bugun, a.tur, a.evet)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    say = Counter(x["onem"] for x in s["sonuc"].values())
    print(f"[OK] {len(s['sonuc'])}/{len(s['basvurular'])} başvuru · Yüksek {say['Yüksek']} · Orta {say['Orta']} · Düşük {say['Düşük']}")
    for k, (son, durum, _) in s["sureler"].items():
        if durum.startswith(("GECİKTİ", "Süre doluyor", "Geç")):
            print(f"[!] {k}: {durum} (son cevap {son:%d.%m.%Y})")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
