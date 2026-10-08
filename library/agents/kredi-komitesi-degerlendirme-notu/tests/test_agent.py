"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

DOSYA = KLASOR / "ornek_veri" / "ornek_gida"
BUGUN = date(2026, 10, 8)
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return {"talep_ozeti": "[FİRMA] toplam limitin 10.000.000 TL'den 25.500.000 TL'ye çıkarılmasını talep ediyor.",
            "firma_ve_faaliyet": "Bisküvi ve kraker üreticisi; ortaklar [KİŞİ-2] ve [KİŞİ-3].",
            "mali_degerlendirme": "Satışlar büyüyor, FAVÖK marjı %11,2'den %9,8'e geriledi; asgari DSCR 1,12.",
            "teminat_degerlendirmesi": "Mevcut teminatlar %70,6 karşılıyor.",
            "guclu_yonler": [{"aciklama": "Ödemeler genel olarak düzenli.", "kaynak": "istihbarat"}],
            "zayif_yonler": [{"aciklama": "Kârlılık geriliyor.", "kaynak": "rasyolar"}],
            "riskler": [{"aciklama": "Yatırımın gecikmesi.", "onem": "orta", "azaltici": "12 ay ödemesiz dönem", "kaynak": "limitler"}],
            "oneri": {"karar": "Olumlu", "gerekce": "Yatırım satışları 45.000.000 TL artıracak."},
            "sartlar": [{"sart": "İpotek ekspertizi güncellenmeli.", "zaman": "Kullandırım öncesi"}],
            "izleme_kriterleri": ["Asgari DSCR yıllık izlenmeli; eşik komitece belirlenmeli."],
            "eksik_bilgiler": ["Zincir market sözleşmesi."]}


class KontrolTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.g, cls.m, cls.a = agent.dosyalari_oku(DOSYA)
        cls.k = agent.kontrol_et(cls.g, cls.m, D(100000000))
        cls.metin = "\n".join(f"{o} {a}" for o, a, _ in cls.k["bulgular"])

    def test_okuma(self):
        self.assertEqual({d.ad: d.tur for d in self.m}["teklif.txt"], "Teklif bilgisi")
        t = self.k["teklif"]
        self.assertEqual((t["firma"], t["vkn"], t["grup"], t["grup_riski"]),
                         ("Örnek Gıda Sanayi ve Ticaret A.Ş.", "9876543217", "Kurgu Gıda Grubu", D(6500000)))
        self.assertEqual(self.k["toplam"], {"mevcut": D(10000000), "risk": D(8600000), "talep": D(25500000), "nakdi": D(23000000), "gayrinakdi": D(2500000)})
        self.assertEqual((self.k["teminat"]["mevcut"], self.k["teminat"]["tesis"]), (D(18000000), D(9000000)))

    def test_bulgular(self):
        self.assertIn("orta Toplam limit 10.000.000 TL'den 25.500.000 TL'ye çıkıyor (%155,0 artış)", self.metin)
        self.assertIn("orta Mevcut teminatlar talebin %70,6'ini karşılıyor", self.metin)
        self.assertNotIn("Teminat karşılama oranı", self.metin)                 # tesis edilecekler dahil %105,9
        self.assertIn("orta Net finansal borç / FAVÖK: 2,4 → 3,1 (%29,2 kötüleşme)", self.metin)
        self.assertIn("bilgi Cari oran: 1,35 → 1,18", self.metin)
        self.assertIn("orta Asgari DSCR (yatırım dahil) 1,12: borç servis karşılama sınırda", self.metin)
        self.assertNotIn("Özkaynak", self.metin)                               # iyileşme
        self.assertIn("yüksek Risk grubu toplamı (32.000.000 TL) banka özkaynağının %32,0'i", self.metin)
        self.assertIn("orta 'gecikme': Örnek Bank B'de", self.metin)
        self.assertNotIn("protesto ve icra", self.metin)                       # "bulunmamaktadır"
        self.assertNotIn("Teklifte", self.metin)

    def test_bk54_sinir_ici(self):
        k = agent.kontrol_et(self.g, self.m, D(200000000))
        self.assertFalse(any("BK md. 54" in kay for _, _, kay in k["bulgular"]))
        self.assertFalse(any("BK md. 54" in kay for _, _, kay in agent.kontrol_et(self.g, self.m)["bulgular"]))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_uctan_uca(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "k.xlsx"
            s = agent.calistir(DOSYA, cikti, ["Hasan Örnekoğlu", "Elif Örnekoğlu"], BUGUN, evet=True, banka_ozkaynak=D(100000000))
            m = MESAJLAR[0]
            for gizli in ("Örnek Gıda", "9876543217", "Kurgu Gıda", "Deniz Örnekçi", "Hasan", "Örnekoğlu"):
                self.assertNotIn(gizli, m)
            for var in ("[FİRMA]", "[GRUP]", "[KİŞİ-2]", "<kod_kontrolleri>", "Makine rehni"):
                self.assertIn(var, m)
            y = s["yanit"]
            self.assertEqual(y["dogrulanamayan_sayilar"], ["45000000"])            # kaynakta 40.000.000
            capraz = " ".join(a for _, a, _ in y["capraz"])
            self.assertIn("Model önerisi 'Olumlu' ama kodda", capraz)
            self.assertIn("Tesis edilecek teminat 'Makine rehni (yeni üretim hattı)' şartlarda yer almıyor", capraz)
            md = s["md"].read_text(encoding="utf-8")
            self.assertIn("ortaklar Hasan Örnekoğlu ve Elif Örnekoğlu", md)
            self.assertIn("Örnek Gıda Sanayi ve Ticaret A.Ş. toplam limitin", md)
            self.assertIn("| **Toplam** | **10.000.000** | **8.600.000** | **25.500.000** |", md)
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Özet", "Kontroller", "Limitler", "Teminatlar", "Rasyolar", "Şartlar ve İzleme"])
            self.assertEqual(wb["Limitler"]["F6"].value, 25500000)
            self.assertEqual(wb["Rasyolar"]["E8"].value, "Kötüleşme")              # Kaldıraç

    def test_klasor_yok(self):
        with self.assertRaises(llm.LLMHatasi):
            agent.calistir(Path("olmayan_klasor"), Path("x.xlsx"), evet=True)


if __name__ == "__main__":
    unittest.main()
