"""
Kredi Komitesi Değerlendirme Notu — Workers / Workless AI Agent
Bankacılık › Krediler Tahsis › Kredi Tahsis Uzmanı

1. Teklif klasöründeki dosyalar okunur: teklif bilgisi (.txt: Firma, VKN, Sektör, Risk Grubu, Grup Riski, Talep
   Gerekçesi, Ödeme Kaynağı), limit tablosu (mevcut / talep), teminat tablosu, rasyo tablosu (iki yıl veya daha
   fazla) ve istihbarat, analiz, ziyaret notları gibi metinler (.txt/.pdf/.docx).
2. Kod hesaplar ve kontrol eder: mevcut ve talep edilen limitler (nakdi / gayrinakdi), limit artışı, limit
   doluluğu, teminat karşılama oranı (mevcut / tesis edilecek ayrı), rasyolardaki kötüleşme ve eşik altı değerler
   (DSCR, cari oran, faiz karşılama, özkaynak), isteğe bağlı olarak Bankacılık Kanunu 54. madde %25 kredi sınırı ve
   metinlerdeki olumsuz ifadeler.
3. Firma unvanı, VKN, risk grubu, müşteri temsilcisi ve --gizle ile verilen adlar takma adlarla; telefon, e-posta,
   TCKN ve IBAN maskelenir. Model; güçlü ve zayıf yönleri, riskleri ve azaltıcı unsurları, öneriyi (taslak), şartları
   ve izleme kriterlerini yazar. Kod, modelin önerisini kod bulgularıyla, şartları tesis edilecek teminatlarla
   karşılaştırır; girdilerde olmayan sayıları işaretler.
4. Çıktı: Markdown komite notu (tablolar koddan) + Excel (Özet + 'Komite Kararı', Kontroller + 'İnceleme',
   Limitler, Teminatlar, Rasyolar, Şartlar ve İzleme).

Kullanım:
    python agent.py                                          # örnek: kurgusal gıda üreticisi
    python agent.py --girdi ./teklif_klasoru --gizle "Ortak Adı" "Kefil Adı"
    python agent.py --girdi ./teklif --banka-ozkaynak 25000000000
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
SIFIR = Decimal(0)
OLUMSUZ = [("karşılıksız", "yüksek"), ("protest", "yüksek"), ("haciz", "yüksek"), ("icra", "yüksek"), ("iflas", "yüksek"),
           ("konkordato", "yüksek"), ("yasal takip", "yüksek"), ("takibe", "yüksek"), ("gecikme", "orta"), ("sark", "orta"),
           ("yapılandır", "orta"), ("dava", "orta"), ("olumsuz", "orta"), ("geç öde", "orta")]
OLUMSUZLAMA = re.compile(r"\b(yok|yoktur|bulunmamaktadır|bulunmamakta|rastlanmamıştır|görülmemiştir|duyulmadı|olmamıştır|yaşanmamıştır)\b")
GAYRINAKDI = ("mektub", "mektup", "akreditif", "gayrinakdi", "gayri nakdi", "kabul kredisi", "aval", "garanti")
# (anahtar, iyi yön: +1 yüksek iyi / -1 düşük iyi / 0 yön yok, eşik kontrolü)
RASYOLAR = [("dscr", 1, "dscr"), ("borc servis", 1, "dscr"), ("net finansal borc", -1, None), ("nfb", -1, None),
            ("faiz karsilama", 1, "faiz"), ("cari oran", 1, "cari"), ("asit", 1, None), ("likidite", 1, None),
            ("kaldirac", -1, None), ("borc ozkaynak", -1, None), ("devir suresi", -1, None), ("tahsil suresi", -1, None),
            ("marj", 1, None), ("buyume", 0, "buyume"), ("ozkaynak", 1, "ozkaynak")]
BK54_ORAN = Decimal("0.25")   # 5411 s. Bankacılık Kanunu md. 54: bir kişi veya risk grubuna kredi toplamı özkaynağın %25'ini aşamaz


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def katla(s) -> str:
    s = kucuk(s)
    for a, b in zip("ıöüşğçâîû", "iousgcaiu"):
        s = s.replace(a, b)
    return re.sub(r"[^a-z0-9%]+", " ", s).strip()


def para(x) -> Decimal | None:
    if x in (None, "", "—", "-"):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("₺", "").replace("%", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tl(x) -> str:
    return "—" if x is None else f"{x:,.0f}".replace(",", ".")


def yuzde(x) -> str:
    return "—" if x is None else f"%{x * 100:.1f}".replace(".", ",")


def sayi(x, yuzdeli=False) -> str:
    if x is None:
        return "—"
    if abs(x) >= 1000:
        return tl(x)
    s = f"{x:.2f}".rstrip("0").rstrip(".").replace(".", ",")
    return f"%{s}" if yuzdeli else s


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(str(x or "").strip(), f).date()
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = belge.txt_oku(yol)
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def basliklar(s: list[list], *anahtarlar: str) -> int:
    """Başlık satırının indeksi: anahtarlardan birini içeren ilk satır (üstteki başlık satırları atlanır)."""
    return next((i for i, r in enumerate(s[:10]) if any(katla(c).startswith(anahtarlar) for c in r if c)), 0)


def sutun(b: list[str], *adlar: str) -> int | None:
    for a in adlar:
        for j, x in enumerate(b):
            if x == a or x.startswith(a):
                return j
    return None


# ----------------------------------------------------------------------------
# Girdiler
# ----------------------------------------------------------------------------

@dataclass
class Dosya:
    ad: str
    tur: str
    metin: str


def limit_oku(yol: Path) -> list[dict]:
    s = tablo_oku(yol)
    bi = basliklar(s, "kredi", "limit turu", "urun")
    b = [katla(c) for c in s[bi]]
    it, iml, imr = sutun(b, "kredi", "limit turu", "urun", "tur"), sutun(b, "mevcut limit"), sutun(b, "mevcut risk", "risk")
    itl, iv, if_ = sutun(b, "talep", "teklif edilen", "onerilen"), sutun(b, "vade"), sutun(b, "fiyat", "faiz", "komisyon")
    al = lambda r, j: r[j] if j is not None and j < len(r) else None  # noqa: E731
    sonuc = []
    for r in s[bi + 1:]:
        ad = str(al(r, it) or "").strip()
        if not ad or katla(ad).startswith(("toplam", "genel")):
            continue
        sonuc.append({"tur": ad, "mevcut": para(al(r, iml)) or SIFIR, "risk": para(al(r, imr)) or SIFIR, "talep": para(al(r, itl)) or SIFIR,
                      "vade": str(al(r, iv) or "").strip(), "fiyat": str(al(r, if_) or "").strip(),
                      "gayrinakdi": any(g in katla(ad) for g in map(katla, GAYRINAKDI))})
    return sonuc


def teminat_oku(yol: Path) -> list[dict]:
    s = tablo_oku(yol)
    bi = basliklar(s, "teminat")
    b = [katla(c) for c in s[bi]]
    it, id_, idu, ia = sutun(b, "teminat", "tur"), sutun(b, "deger", "tutar"), sutun(b, "durum"), sutun(b, "aciklama")
    al = lambda r, j: r[j] if j is not None and j < len(r) else None  # noqa: E731
    sonuc = []
    for r in s[bi + 1:]:
        ad = str(al(r, it) or "").strip()
        if not ad or katla(ad).startswith("toplam"):
            continue
        durum = katla(al(r, idu))
        sonuc.append({"tur": ad, "deger": para(al(r, id_)), "tesis": "tesis" in durum or "alinacak" in durum or "edilecek" in durum,
                      "durum": str(al(r, idu) or "").strip(), "aciklama": str(al(r, ia) or "").strip()})
    return sonuc


def rasyo_oku(yol: Path) -> tuple[list[str], list[dict]]:
    s = tablo_oku(yol)
    bi = basliklar(s, "rasyo", "oran", "gosterge", "kalem")
    donemler = [str(c).strip() for c in s[bi][1:] if c not in (None, "")]
    sonuc = []
    for r in s[bi + 1:]:
        if not r or not r[0]:
            continue
        ham = [r[j] if j < len(r) else None for j in range(1, len(donemler) + 1)]
        degerler = [para(x) for x in ham]
        yuzdeli = any("%" in str(x or "") for x in ham)
        k = katla(r[0])
        yon, esik = next(((y, e) for a, y, e in RASYOLAR if a in k), (0, None))
        sonuc.append({"ad": str(r[0]).strip(), "degerler": degerler, "yuzdeli": yuzdeli, "yon": yon, "esik": esik})
    return donemler, sonuc


def anahtarli(metin: str, *etiketler: str) -> str | None:
    for et in etiketler:
        m = re.search(r"^\s*" + re.escape(et) + r"(?:\s*\([^)\n]*\))?\s*:\s*(.+)$", metin or "", re.I | re.M)
        if m:
            return m.group(1).strip()
    return None


def dosyalari_oku(klasor: Path) -> tuple[dict, list[Dosya], list[str]]:
    girdi = {"teklif": None, "limitler": [], "teminatlar": [], "rasyolar": [], "donemler": []}
    metinler, atlanan = [], []
    for yol in sorted(klasor.iterdir()):
        ad = katla(yol.stem)
        if yol.suffix.lower() in {".csv", ".xlsx", ".xlsm"}:
            try:
                ilk = katla(" ".join(str(c) for r in tablo_oku(yol)[:3] for c in r if c))
            except Exception as h:  # noqa: BLE001
                atlanan.append(f"{yol.name} ({h})")
                continue
            if "limit" in ad or "talep" in ilk:
                girdi["limitler"], tur = limit_oku(yol), "Limit tablosu"
            elif "teminat" in ad or ilk.startswith("teminat"):
                girdi["teminatlar"], tur = teminat_oku(yol), "Teminat tablosu"
            elif "rasyo" in ad or "oran" in ad or ilk.startswith(("rasyo", "oran", "gosterge")):
                girdi["donemler"], girdi["rasyolar"] = rasyo_oku(yol)
                tur = "Rasyo tablosu"
            else:
                atlanan.append(f"{yol.name} (tanınmayan tablo)")
                continue
            metinler.append(Dosya(yol.name, tur, ""))
            continue
        if yol.suffix.lower() not in belge.DESTEKLENEN:
            atlanan.append(yol.name)
            continue
        try:
            metin = belge.metin_oku(yol)
        except belge.BelgeHatasi as h:
            atlanan.append(f"{yol.name} ({h})")
            continue
        bas = katla(metin[:400])
        if girdi["teklif"] is None and "teklif" in ad:
            girdi["teklif"], tur = metin, "Teklif bilgisi"
        elif "istihbarat" in ad or "istihbarat" in bas:
            tur = "İstihbarat"
        elif "analiz" in ad or "analist" in ad or "analiz" in bas:
            tur = "Analiz notu"
        elif "ziyaret" in ad or "ziyaret" in bas:
            tur = "Ziyaret notu"
        elif girdi["teklif"] is None and anahtarli(metin, "Firma", "Unvan") and anahtarli(metin, "Talep Gerekçesi", "Talep"):
            girdi["teklif"], tur = metin, "Teklif bilgisi"
        else:
            tur = "Diğer not"
        metinler.append(Dosya(yol.name, tur, metin))
    return girdi, metinler, atlanan


# ----------------------------------------------------------------------------
# Kontroller
# ----------------------------------------------------------------------------

def kontrol_et(girdi: dict, metinler: list[Dosya], banka_ozkaynak: Decimal | None = None) -> dict:
    bulgular = []

    def b(onem, aciklama, kaynak=""):
        bulgular.append((onem, aciklama, kaynak))

    t = girdi["teklif"] or ""
    teklif = {"firma": anahtarli(t, "Firma", "Unvan", "Müşteri"), "vkn": re.sub(r"\D", "", anahtarli(t, "VKN", "Vergi") or ""),
              "sektor": anahtarli(t, "Sektör", "Faaliyet"), "grup": anahtarli(t, "Risk Grubu", "Grup"),
              "grup_riski": para(re.sub(r"[^\d.,]", "", anahtarli(t, "Grup Riski") or "")),
              "temsilci": anahtarli(t, "Müşteri Temsilcisi", "Portföy Yöneticisi", "Temsilci")}
    if not girdi["teklif"]:
        b("orta", "Teklif bilgisi dosyası (teklif.txt) yok: firma, talep gerekçesi ve ödeme kaynağı bilinmiyor")
    for et in ("Talep Gerekçesi", "Ödeme Kaynağı"):
        if girdi["teklif"] and not anahtarli(t, et):
            b("orta", f"Teklifte '{et}' yazılmamış", "teklif")

    lim = girdi["limitler"]
    top = {a: sum((x[a] for x in lim), SIFIR) for a in ("mevcut", "risk", "talep")}
    top["nakdi"] = sum((x["talep"] for x in lim if not x["gayrinakdi"]), SIFIR)
    top["gayrinakdi"] = top["talep"] - top["nakdi"]
    if not lim:
        b("yüksek", "Limit tablosu yok: talep edilen limitler bilinmiyor")
    else:
        if top["mevcut"] and top["talep"] / top["mevcut"] >= Decimal("1.5"):
            b("orta", f"Toplam limit {tl(top['mevcut'])} TL'den {tl(top['talep'])} TL'ye çıkıyor ({yuzde(top['talep'] / top['mevcut'] - 1)} artış)", "limitler")
        for x in lim:
            if not x["mevcut"] and x["talep"]:
                b("bilgi", f"Yeni kredi türü: {x['tur']} {tl(x['talep'])} TL" + (f", vade {x['vade']}" if x["vade"] else ""), "limitler")
            elif x["mevcut"] and x["risk"] / x["mevcut"] >= Decimal("0.9"):
                b("bilgi", f"{x['tur']}: mevcut limit doluluğu {yuzde(x['risk'] / x['mevcut'])} ({tl(x['risk'])} / {tl(x['mevcut'])} TL)", "limitler")
            if x["risk"] > x["mevcut"] and x["mevcut"]:
                b("yüksek", f"{x['tur']}: risk ({tl(x['risk'])} TL) mevcut limiti ({tl(x['mevcut'])} TL) aşıyor", "limitler")

    tem = girdi["teminatlar"]
    tm = {"mevcut": sum((x["deger"] or SIFIR for x in tem if not x["tesis"]), SIFIR), "tesis": sum((x["deger"] or SIFIR for x in tem if x["tesis"]), SIFIR)}
    tm["toplam"] = tm["mevcut"] + tm["tesis"]
    tm["oran_mevcut"] = tm["mevcut"] / top["talep"] if top["talep"] else None
    tm["oran_toplam"] = tm["toplam"] / top["talep"] if top["talep"] else None
    if not tem:
        b("orta", "Teminat tablosu yok", "teminatlar")
    elif top["talep"]:
        if tm["oran_toplam"] < 1:
            b("orta", f"Teminat karşılama oranı {yuzde(tm['oran_toplam'])} ({tl(tm['toplam'])} / {tl(top['talep'])} TL): teminat talep edilen limitin altında", "teminatlar")
        if tm["tesis"] and tm["oran_mevcut"] < 1:
            b("orta", f"Mevcut teminatlar talebin {yuzde(tm['oran_mevcut'])}'ini karşılıyor; tesis edilecek teminatlar ({tl(tm['tesis'])} TL) "
                      "olmadan karşılama yetersiz — kullandırım şartına bağlanmalı", "teminatlar")
        for x in tem:
            if x["deger"] is None:
                b("bilgi", f"{x['tur']}: değer yazılmamış, karşılama oranına katılmadı", "teminatlar")
        b("bilgi", "Teminat değerleri ham değerlerdir; bankanızın teminat katsayıları (ör. ipotekte ekspertiz değerinin belli bir oranı) uygulanmadı", "teminatlar")

    for r in girdi["rasyolar"]:
        d = [x for x in r["degerler"] if x is not None]
        if not d:
            continue
        son = r["degerler"][-1] if r["degerler"][-1] is not None else d[-1]
        oncekiler = [x for x in r["degerler"][:-1] if x is not None]
        r["son"], r["onceki"] = son, (oncekiler[-1] if oncekiler else None)
        r["degisim"] = r["degerlendirme"] = None
        if r["onceki"] not in (None, SIFIR) and r["yon"]:
            r["degisim"] = (son - r["onceki"]) / abs(r["onceki"])
            kotu = -r["yon"] * r["degisim"]
            if kotu >= Decimal("0.10"):
                r["degerlendirme"] = "Kötüleşme"
                b("orta" if kotu >= Decimal("0.25") else "bilgi",
                  f"{r['ad']}: {sayi(r['onceki'], r['yuzdeli'])} → {sayi(son, r['yuzdeli'])} ({yuzde(kotu)} kötüleşme)", "rasyolar")
            elif -kotu >= Decimal("0.10"):
                r["degerlendirme"] = "İyileşme"
        e = r["esik"]
        if e == "dscr" and son < 1:
            b("yüksek", f"{r['ad']} {sayi(son)}: nakit akışı borç servisini karşılamıyor (< 1)", "rasyolar")
        elif e == "dscr" and son < Decimal("1.2"):
            b("orta", f"{r['ad']} {sayi(son)}: borç servis karşılama sınırda (1 – 1,2)", "rasyolar")
        elif e == "cari" and son < 1:
            b("orta", f"{r['ad']} {sayi(son)}: kısa vadeli borçlar dönen varlıklardan fazla", "rasyolar")
        elif e == "faiz" and son < 1:
            b("yüksek", f"{r['ad']} {sayi(son)}: faaliyet kârı faiz giderini karşılamıyor", "rasyolar")
        elif e == "ozkaynak" and son <= 0:
            b("yüksek", f"{r['ad']} {tl(son)} TL: özkaynak negatif", "rasyolar")
        elif e == "buyume" and son < 0:
            b("bilgi", f"{r['ad']} {sayi(son, r['yuzdeli'])}: satışlar daralmış", "rasyolar")
    if not girdi["rasyolar"]:
        b("orta", "Rasyo tablosu yok: mali değerlendirme yalnız metinlere dayanacak")

    bk54 = None
    if banka_ozkaynak:
        toplam = top["talep"] + (teklif["grup_riski"] or SIFIR)
        bk54 = toplam / banka_ozkaynak
        if bk54 > BK54_ORAN:
            b("yüksek", f"Risk grubu toplamı ({tl(toplam)} TL) banka özkaynağının {yuzde(bk54)}'i: Bankacılık Kanunu md. 54 %25 sınırını aşıyor", "BK md. 54")
        elif bk54 > Decimal("0.20"):
            b("bilgi", f"Risk grubu toplamı banka özkaynağının {yuzde(bk54)}'i: %25 sınırına yakın", "BK md. 54")
    for d in metinler:
        if d.tur in ("Teklif bilgisi",) or not d.metin:
            continue
        for cumle in re.split(r"(?<=[.!?])\s+|\n", d.metin):
            ck = kucuk(cumle)
            for kelime, onem in OLUMSUZ:
                if kelime in ck and not OLUMSUZLAMA.search(ck):
                    b(onem, f"'{kelime}': {cumle.strip()[:170]}", d.ad)
                    break
    return {"bulgular": bulgular, "teklif": teklif, "limitler": lim, "toplam": top, "teminatlar": tem, "teminat": tm,
            "rasyolar": girdi["rasyolar"], "donemler": girdi["donemler"], "bk54": bk54}


# ----------------------------------------------------------------------------
# Maskeleme
# ----------------------------------------------------------------------------

def takma_adlar(teklif: dict, terimler: list[str]) -> dict[str, str]:
    harita = {}
    if teklif.get("firma"):
        harita[teklif["firma"]] = "[FİRMA]"
        kok = re.split(r"\s+(?:San\.|Sanayi|Tic\.|Ticaret|A\.Ş\.|Ltd\.)", teklif["firma"])[0].strip()
        if len(kok) > 4 and kok != teklif["firma"]:
            harita[kok] = "[FİRMA]"
    if teklif.get("vkn"):
        harita[teklif["vkn"]] = "[VKN]"
    if teklif.get("grup"):
        harita[teklif["grup"]] = "[GRUP]"
    kisiler = [teklif["temsilci"]] if teklif.get("temsilci") else []
    for i, ad in enumerate(dict.fromkeys(a.strip() for a in kisiler + list(terimler) if a.strip()), 1):
        harita[ad] = f"[KİŞİ-{i}]"
        parcalar = ad.split()
        if len(parcalar) > 1 and len(parcalar[-1]) > 3:
            harita.setdefault(parcalar[-1], f"[KİŞİ-{i}-SOYAD]")
    return harita


def maskele(metin: str, harita: dict[str, str]) -> str:
    for gercek, takma in sorted(harita.items(), key=lambda x: -len(x[0])):
        metin = re.sub(re.escape(gercek), takma, metin, flags=re.I)
    return llm.maskele(metin)


def geri_ac(metin: str, harita: dict[str, str]) -> str:
    for gercek, takma in sorted(harita.items(), key=lambda x: -len(x[1])):
        if takma != "[FİRMA]":
            metin = metin.replace(takma, gercek)
    unvan = next((g for g, t in harita.items() if t == "[FİRMA]"), None)
    return metin.replace("[FİRMA]", unvan) if unvan else metin


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

KAYNAKLI = {"type": "object", "properties": {"aciklama": {"type": "string"}, "kaynak": {"type": "string"}},
            "required": ["aciklama", "kaynak"], "additionalProperties": False}
KARARLAR = ["Olumlu", "Şartlı olumlu", "Ek bilgi gerekli", "Olumsuz"]
SEMA = {
    "type": "object",
    "properties": {
        "talep_ozeti": {"type": "string"},
        "firma_ve_faaliyet": {"type": "string"},
        "mali_degerlendirme": {"type": "string"},
        "teminat_degerlendirmesi": {"type": "string"},
        "guclu_yonler": {"type": "array", "items": KAYNAKLI},
        "zayif_yonler": {"type": "array", "items": KAYNAKLI},
        "riskler": {"type": "array", "items": {
            "type": "object", "properties": {"aciklama": {"type": "string"}, "onem": {"type": "string", "enum": ["yüksek", "orta", "düşük"]},
                                             "azaltici": {"type": "string"}, "kaynak": {"type": "string"}},
            "required": ["aciklama", "onem", "azaltici", "kaynak"], "additionalProperties": False}},
        "oneri": {"type": "object", "properties": {"karar": {"type": "string", "enum": KARARLAR}, "gerekce": {"type": "string"}},
                  "required": ["karar", "gerekce"], "additionalProperties": False},
        "sartlar": {"type": "array", "items": {
            "type": "object", "properties": {"sart": {"type": "string"},
                                             "zaman": {"type": "string", "enum": ["Kullandırım öncesi", "Kullandırım sonrası", "Sürekli"]}},
            "required": ["sart", "zaman"], "additionalProperties": False}},
        "izleme_kriterleri": {"type": "array", "items": {"type": "string"}},
        "eksik_bilgiler": {"type": "array", "items": {"type": "string"}},
    },
    "required": ["talep_ozeti", "firma_ve_faaliyet", "mali_degerlendirme", "teminat_degerlendirmesi", "guclu_yonler", "zayif_yonler",
                 "riskler", "oneri", "sartlar", "izleme_kriterleri", "eksik_bilgiler"],
    "additionalProperties": False,
}
METIN_ALANLARI = ("talep_ozeti", "firma_ve_faaliyet", "mali_degerlendirme", "teminat_degerlendirmesi")


def kod_ozeti(k: dict) -> list[str]:
    t, top, tm = k["teklif"], k["toplam"], k["teminat"]
    satir = [f"- Firma: {t['firma'] or '—'}; sektör: {t['sektor'] or '—'}; risk grubu: {t['grup'] or '—'}"]
    if t["grup_riski"]:
        satir.append(f"- Grup riski (bankamız, bu firma hariç): {tl(t['grup_riski'])} TL")
    for x in k["limitler"]:
        satir.append(f"- Limit {x['tur']} ({'gayrinakdi' if x['gayrinakdi'] else 'nakdi'}): mevcut {tl(x['mevcut'])} TL, risk {tl(x['risk'])} TL, "
                     f"talep {tl(x['talep'])} TL" + (f", vade {x['vade']}" if x["vade"] else "") + (f", fiyat {x['fiyat']}" if x["fiyat"] else ""))
    if k["limitler"]:
        satir.append(f"- Limit toplamı: mevcut {tl(top['mevcut'])} TL, risk {tl(top['risk'])} TL, talep {tl(top['talep'])} TL "
                     f"(nakdi {tl(top['nakdi'])}, gayrinakdi {tl(top['gayrinakdi'])})")
    for x in k["teminatlar"]:
        satir.append(f"- Teminat {x['tur']}: {tl(x['deger'])} TL, {x['durum'] or '—'}" + (f" ({x['aciklama']})" if x["aciklama"] else ""))
    if k["teminatlar"] and top["talep"]:
        satir.append(f"- Teminat karşılama: mevcut {yuzde(tm['oran_mevcut'])}, tesis edilecekler dahil {yuzde(tm['oran_toplam'])} (ham değer)")
    for r in k["rasyolar"]:
        if "son" in r:
            satir.append(f"- Rasyo {r['ad']}: " + ", ".join(f"{d} {sayi(v, r['yuzdeli'])}" for d, v in zip(k["donemler"], r["degerler"]) if v is not None))
    if k["bk54"] is not None:
        satir.append(f"- Risk grubu toplamının banka özkaynağına oranı: {yuzde(k['bk54'])} (BK md. 54 sınırı %25)")
    return satir


def sayilar(metin: str) -> set[str]:
    """3+ haneli sayıların yalnız rakamları ('15.000.000' → '15000000')."""
    return {re.sub(r"\D", "", m) for m in re.findall(r"\d[\d.,]*\d", metin) if len(re.sub(r"\D", "", m)) >= 3}


def hazirla(metinler: list[Dosya], k: dict, harita: dict) -> tuple[str, str]:
    kk = [f"- [{o}] {maskele(a, harita)} (kaynak: {kay or '-'})" for o, a, kay in k["bulgular"]]
    mesaj = "\n".join([
        "<belgeler>",
        *[f'<belge ad="{d.ad}" tur="{d.tur}">\n{maskele(d.metin, harita)}\n</belge>' for d in metinler if d.metin],
        "</belgeler>",
        "<kod_ozeti>", *[maskele(x, harita) for x in kod_ozeti(k)], "</kod_ozeti>",
        "<kod_kontrolleri>", *(kk or ["- Kontrol bulgusu yok"]), "</kod_kontrolleri>",
    ])
    return (BURASI / "prompt.md").read_text(encoding="utf-8"), mesaj


def model_sonrasi(y: dict, k: dict) -> list[tuple[str, str, str]]:
    """Modelin önerisini ve şartlarını kod bulgularıyla karşılaştırır."""
    ek = []
    yuksek = [a for o, a, _ in k["bulgular"] if o == "yüksek"]
    if y["oneri"]["karar"] == "Olumlu" and yuksek:
        ek.append(("yüksek", f"Model önerisi 'Olumlu' ama kodda {len(yuksek)} yüksek önemli bulgu var: {yuksek[0][:100]}", "model ↔ kod"))
    sartlar = katla(" ".join(s["sart"] for s in y["sartlar"]))
    for x in k["teminatlar"]:
        if x["tesis"]:
            kokler = [w[:5] for w in katla(x["tur"]).split() if len(w) >= 5][:2]
            if kokler and not any(kk in sartlar for kk in kokler):
                ek.append(("orta", f"Tesis edilecek teminat '{x['tur']}' şartlarda yer almıyor", "model ↔ teminatlar"))
    return ek


def model_yaz(metinler, k, harita) -> dict:
    sistem, mesaj = hazirla(metinler, k, harita)
    yanit = llm.json_iste(sistem, mesaj, SEMA)
    bilinen = sayilar(mesaj) | {x for x in re.findall(r"\d+", mesaj) if len(x) >= 3}
    metin = " ".join([yanit[a] for a in METIN_ALANLARI] + [yanit["oneri"]["gerekce"]] + yanit["izleme_kriterleri"] + yanit["eksik_bilgiler"] +
                     [s["sart"] for s in yanit["sartlar"]] + [x["aciklama"] for a in ("guclu_yonler", "zayif_yonler", "riskler") for x in yanit[a]] +
                     [x["azaltici"] for x in yanit["riskler"]])
    yanit["dogrulanamayan_sayilar"] = sorted(s for s in sayilar(metin) if s not in bilinen)
    yanit["capraz"] = model_sonrasi(yanit, k)
    return yanit


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
ONEM = {"yüksek": "FDE2E1", "orta": "FFF4CE", "bilgi": "E8F0FE", "düşük": "E8F0FE"}
KARAR = PatternFill("solid", fgColor="FFF4CE")
MODEL = PatternFill("solid", fgColor="EDE7FF")
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, basliklar_, genislik):
    ws.append(basliklar_)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def rapor_yaz(cikti: Path, metinler, atlanan, k: dict, y: dict, harita: dict, bugun: date) -> Path:
    g = lambda s: geri_ac(s, harita)  # noqa: E731
    t, top, tm = k["teklif"], k["toplam"], k["teminat"]
    tum_bulgular = k["bulgular"] + y["capraz"]
    wb = Workbook()
    o = wb.active
    o.title = "Özet"
    o.column_dimensions["A"].width, o.column_dimensions["B"].width = 32, 100
    for a, d in [("Firma", t["firma"] or "—"), ("Sektör", t["sektor"] or "—"), ("Not tarihi", f"{bugun:%d.%m.%Y}"),
                 ("Talep toplamı", f"{tl(top['talep'])} TL (mevcut {tl(top['mevcut'])} TL)"),
                 ("Teminat karşılama", f"mevcut {yuzde(tm['oran_mevcut'])} · tesis edilecekler dahil {yuzde(tm['oran_toplam'])}"),
                 ("Kaynaklar", "; ".join(f"{d.ad} ({d.tur})" for d in metinler)), ("Okunamayan", "; ".join(atlanan) or "—"),
                 ("Talep özeti (model)", g(y["talep_ozeti"])), ("Firma ve faaliyet (model)", g(y["firma_ve_faaliyet"])),
                 ("Mali değerlendirme (model)", g(y["mali_degerlendirme"])), ("Teminat değerlendirmesi (model)", g(y["teminat_degerlendirmesi"])),
                 ("Öneri (taslak, model)", f"{y['oneri']['karar']} — {g(y['oneri']['gerekce'])}"),
                 ("Doğrulanamayan sayılar", ", ".join(y["dogrulanamayan_sayilar"]) or "—"), ("Model", llm.kullanim_ozeti()),
                 ("Komite Kararı", ""), ("Karar Gerekçesi / Notlar", "")]:
        o.append([a, d])
        o.cell(o.max_row, 1).font = Font(bold=True)
        if "model" in a:
            o.cell(o.max_row, 2).fill = MODEL
        if a.startswith(("Komite", "Karar")):
            o.cell(o.max_row, 2).fill = KARAR
        if a == "Doğrulanamayan sayılar" and y["dogrulanamayan_sayilar"]:
            o.cell(o.max_row, 2).fill = PatternFill("solid", fgColor=ONEM["yüksek"])
        for h in o[o.max_row]:
            h.alignment = UST

    kt = wb.create_sheet("Kontroller")
    _baslik(kt, ["Kaynak", "Önem", "Bulgu / Risk", "Dayanak", "Azaltıcı Unsur", "İnceleme"], (8, 9, 80, 22, 50, 26))
    for on, a, kay in tum_bulgular:
        kt.append(["Kod", on, a, kay, "", ""])
        kt.cell(kt.max_row, 2).fill = PatternFill("solid", fgColor=ONEM[on])
    for x in y["riskler"]:
        kt.append(["Model", x["onem"], g(x["aciklama"]), x["kaynak"], g(x["azaltici"]), ""])
        kt.cell(kt.max_row, 2).fill = PatternFill("solid", fgColor=ONEM[x["onem"]])
    for x in y["zayif_yonler"]:
        kt.append(["Model", "zayıf yön", g(x["aciklama"]), x["kaynak"], "", ""])
    for x in y["guclu_yonler"]:
        kt.append(["Model", "güçlü yön", g(x["aciklama"]), x["kaynak"], "", ""])
    for x in y["eksik_bilgiler"]:
        kt.append(["Model", "eksik bilgi", g(x), "", "", ""])
    for r in kt.iter_rows(min_row=2):
        r[5].fill = KARAR
        for h in r:
            h.alignment = UST

    li = wb.create_sheet("Limitler")
    _baslik(li, ["Kredi Türü", "Nakdi / Gayrinakdi", "Mevcut Limit", "Mevcut Risk", "Doluluk", "Talep Edilen", "Değişim", "Vade", "Fiyat"],
            (30, 14, 15, 15, 9, 15, 15, 22, 24))
    for x in k["limitler"] + ([{"tur": "TOPLAM", "gayrinakdi": None, "vade": "", "fiyat": "", **{a: top[a] for a in ("mevcut", "risk", "talep")}}] if k["limitler"] else []):
        li.append([x["tur"], "" if x["gayrinakdi"] is None else ("Gayrinakdi" if x["gayrinakdi"] else "Nakdi"), float(x["mevcut"]), float(x["risk"]),
                   float(x["risk"] / x["mevcut"]) if x["mevcut"] else None, float(x["talep"]), float(x["talep"] - x["mevcut"]), x["vade"], x["fiyat"]])
        for c in (3, 4, 6, 7):
            li.cell(li.max_row, c).number_format = "#,##0"
        li.cell(li.max_row, 5).number_format = "0.0%"
        if x["tur"] == "TOPLAM":
            for h in li[li.max_row]:
                h.font = Font(bold=True)

    te = wb.create_sheet("Teminatlar")
    _baslik(te, ["Teminat", "Değer (ham)", "Durum", "Açıklama"], (40, 16, 16, 60))
    for x in k["teminatlar"]:
        te.append([x["tur"], x["deger"] and float(x["deger"]), x["durum"], x["aciklama"]])
        te.cell(te.max_row, 2).number_format = "#,##0"
        if x["tesis"]:
            te.cell(te.max_row, 3).fill = PatternFill("solid", fgColor=ONEM["orta"])
    if k["teminatlar"]:
        for a, v in [("Mevcut teminat toplamı", tm["mevcut"]), ("Tesis edilecek", tm["tesis"]), ("Toplam", tm["toplam"]), ("Talep edilen limit", top["talep"])]:
            te.append([a, float(v)])
            te.cell(te.max_row, 1).font = Font(bold=True)
            te.cell(te.max_row, 2).number_format = "#,##0"

    ra = wb.create_sheet("Rasyolar")
    _baslik(ra, ["Rasyo", *k["donemler"], "Değişim", "Değerlendirme"], (40, *[12] * len(k["donemler"]), 11, 14))
    for r in k["rasyolar"]:
        ra.append([r["ad"], *[(float(v) / 100 if r["yuzdeli"] else float(v)) if v is not None else None for v in r["degerler"]],
                   float(r["degisim"]) if r.get("degisim") is not None else None, r.get("degerlendirme") or ""])
        for j in range(len(k["donemler"])):
            ra.cell(ra.max_row, 2 + j).number_format = "0.0%" if r["yuzdeli"] else ("#,##0" if any(v and abs(v) >= 1000 for v in r["degerler"]) else "0.00")
        ra.cell(ra.max_row, 2 + len(k["donemler"])).number_format = "0.0%"
        if r.get("degerlendirme") == "Kötüleşme":
            ra.cell(ra.max_row, 3 + len(k["donemler"])).fill = PatternFill("solid", fgColor=ONEM["orta"])

    sa = wb.create_sheet("Şartlar ve İzleme")
    _baslik(sa, ["Tür", "Zaman", "Şart / Kriter (taslak)", "Komite Kararı"], (14, 18, 90, 20))
    for s in y["sartlar"]:
        sa.append(["Şart", s["zaman"], g(s["sart"]), ""])
    for s in y["izleme_kriterleri"]:
        sa.append(["İzleme", "Sürekli", g(s), ""])
    for r in sa.iter_rows(min_row=2):
        r[3].fill = KARAR
        r[2].alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

    md = cikti.with_suffix(".md")
    sat = [f"# Kredi Komitesi Değerlendirme Notu (taslak) — {t['firma'] or ''}", "",
           f"Not tarihi: {bugun:%d.%m.%Y} · Sektör: {t['sektor'] or '—'} · Risk grubu: {t['grup'] or '—'}", "",
           "## Talep", "", g(y["talep_ozeti"])]
    if k["limitler"]:
        sat += ["", "| Kredi türü | Mevcut limit | Risk | Talep | Vade | Fiyat |", "|---|---:|---:|---:|---|---|"]
        sat += [f"| {x['tur']} | {tl(x['mevcut'])} | {tl(x['risk'])} | {tl(x['talep'])} | {x['vade']} | {x['fiyat']} |" for x in k["limitler"]]
        sat += [f"| **Toplam** | **{tl(top['mevcut'])}** | **{tl(top['risk'])}** | **{tl(top['talep'])}** | | |"]
    sat += ["", "## Firma ve faaliyet", "", g(y["firma_ve_faaliyet"]), "", "## Mali değerlendirme", "", g(y["mali_degerlendirme"])]
    if k["rasyolar"]:
        sat += ["", "| Rasyo | " + " | ".join(k["donemler"]) + " |", "|---|" + "---:|" * len(k["donemler"])]
        sat += [f"| {r['ad']} | " + " | ".join(sayi(v, r["yuzdeli"]) for v in r["degerler"]) + " |" for r in k["rasyolar"]]
    sat += ["", "## Teminatlar", "", g(y["teminat_degerlendirmesi"])]
    if k["teminatlar"]:
        sat += ["", "| Teminat | Değer | Durum |", "|---|---:|---|"] + [f"| {x['tur']} | {tl(x['deger'])} | {x['durum']} |" for x in k["teminatlar"]]
        sat += ["", f"Teminat karşılama (ham değer): mevcut {yuzde(tm['oran_mevcut'])}, tesis edilecekler dahil {yuzde(tm['oran_toplam'])}."]
    sat += ["", "## Güçlü yönler", ""] + [f"- {g(x['aciklama'])} ({x['kaynak']})" for x in y["guclu_yonler"]]
    sat += ["", "## Zayıf yönler", ""] + [f"- {g(x['aciklama'])} ({x['kaynak']})" for x in y["zayif_yonler"]]
    sat += ["", "## Riskler ve azaltıcı unsurlar", "", "| Önem | Risk | Azaltıcı unsur |", "|---|---|---|"]
    sat += [f"| {x['onem']} | {g(x['aciklama'])} | {g(x['azaltici'])} |" for x in y["riskler"]]
    sat += ["", "## Kod kontrolleri", ""] + ([f"- **{o_}** {a} ({kay})" for o_, a, kay in tum_bulgular] or ["- Bulgu yok"])
    sat += ["", "## Öneri (taslak)", "", f"**{y['oneri']['karar']}** — {g(y['oneri']['gerekce'])}"]
    sat += ["", "### Şartlar", ""] + [f"- *{s['zaman']}:* {g(s['sart'])}" for s in y["sartlar"]]
    sat += ["", "### İzleme kriterleri", ""] + [f"- {g(s)}" for s in y["izleme_kriterleri"]]
    if y["eksik_bilgiler"]:
        sat += ["", "## Eksik bilgi / belge", ""] + [f"- {g(s)}" for s in y["eksik_bilgiler"]]
    if y["dogrulanamayan_sayilar"]:
        sat += ["", f"> Girdilerde bulunmayan sayılar (kontrol edin): {', '.join(y['dogrulanamayan_sayilar'])}"]
    sat += ["", "> Taslaktır; kredi kararı yetkili komiteye aittir. Bankacılık sırrı içerir, yalnız yetkili kişilerle paylaşın."]
    md.write_text("\n".join(sat) + "\n", encoding="utf-8")
    return md


def calistir(klasor: Path, cikti: Path, terimler: list[str] | None = None, bugun: date | None = None, evet: bool = False,
             banka_ozkaynak: Decimal | None = None) -> dict:
    if not klasor.is_dir():
        raise llm.LLMHatasi(f"{klasor} klasörü bulunamadı.")
    bugun = bugun or date.today()
    girdi, metinler, atlanan = dosyalari_oku(klasor)
    if not metinler:
        raise llm.LLMHatasi(f"{klasor}: okunabilir dosya yok.")
    k = kontrol_et(girdi, metinler, banka_ozkaynak)
    harita = takma_adlar(k["teklif"], terimler or [])
    for d in metinler:
        print(f"[OK] {d.ad}: {d.tur}")
    kisi = sum(1 for t in harita.values() if t.startswith("[KİŞİ") and not t.endswith("SOYAD]"))
    llm.onay_al(f"{len(metinler)} dosyanın metni ve kod özeti gönderilecek (firma unvanı, VKN, risk grubu ve {kisi} kişi takma adlı; "
                "telefon, e-posta, TCKN, IBAN maskeli).", evet)
    y = model_yaz(metinler, k, harita)
    md = rapor_yaz(cikti, metinler, atlanan, k, y, harita, bugun)
    return {"kontrol": k, "yanit": y, "harita": harita, "md": md, "metinler": metinler}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Teklif, analiz ve istihbarat bulgularından kredi komitesi değerlendirme notu taslağı hazırlar.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "ornek_gida", help="Teklif klasörü")
    p.add_argument("--tarih", help="Not tarihi GG.AA.YYYY (varsayılan: bugün; örnek veride 08.10.2026)")
    p.add_argument("--gizle", nargs="*", default=[], metavar="AD", help="Ayrıca maskelenecek kişi/şirket adları (ortaklar, kefiller)")
    p.add_argument("--banka-ozkaynak", help="Bankanın özkaynağı (TL); verilirse BK md. 54 %%25 kredi sınırı kontrol edilir")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "komite_notu.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    bugun = tarih(a.tarih) if a.tarih else (date(2026, 10, 8) if a.girdi == BURASI / "ornek_veri" / "ornek_gida" else None)
    try:
        s = calistir(a.girdi, a.cikti, a.gizle, bugun, a.evet, para(a.banka_ozkaynak) if a.banka_ozkaynak else None)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    for o, ac, kay in s["kontrol"]["bulgular"] + s["yanit"]["capraz"]:
        if o != "bilgi":
            print(f"[{'X' if o == 'yüksek' else '!'}] {ac}")
    if s["yanit"]["dogrulanamayan_sayilar"]:
        print(f"[!] Girdilerde olmayan sayılar: {', '.join(s['yanit']['dogrulanamayan_sayilar'])}")
    print(f"[i] Öneri (taslak): {s['yanit']['oneri']['karar']}")
    print(f"[OK] Rapor: {a.cikti.resolve()} · {s['md'].resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
