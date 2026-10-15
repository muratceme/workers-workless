"""
Piyasa Günlüğü Özeti — Workers / Workless AI Agent
Hazine › Hazine Uzmanı

1. Kod gün içi piyasa verilerini (Gösterge; Tür; Önceki Kapanış; Son) okur ve değişimleri hesaplar:
   - Kur, endeks, emtia: tutar ve yüzde değişim
   - Faiz (gösterge tahvil, politika faizi): baz puan değişimi; CDS: baz puan değişimi
   - Günün en çok hareket eden fiyat göstergesi ve faiz / CDS göstergesi
   Haber / gelişme notları (isteğe bağlı) H1, H2 … diye numaralanır.
2. Model başlık, öne çıkanlar ve kısa piyasa özeti paragrafları yazar; her madde dayandığı gösterge ve haber numarasını gösterir.
3. Kod denetler:
   - Metindeki her sayı dayandığı göstergenin hazır rakamlarından ya da haber satırından mı
   - Faiz ve CDS değişimi "baz puan" ile mi yazılmış (yüzde ile karıştırılmamış mı)
   - "Yükseldi / geriledi" yönü değişimin işaretiyle tutarlı mı
   - Neden-sonuç ve beklenti ifadeleri bir habere dayanıyor mu; en çok hareket eden göstergeler ele alınmış mı
   - Yatırım tavsiyesi sayılabilecek ifade ("alım fırsatı", "öneririz", "hedef fiyat") var mı
   Sorun varsa modele geri gönderilir (en fazla --tur kez).
4. Metnin sonuna "yatırım danışmanlığı kapsamında değildir" uyarısını kod ekler.

Kullanım:
    python agent.py                                                  # örnek gün verisi
    python agent.py --girdi piyasa.xlsx --haberler haberler.txt --tarih 14.10.2026
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
_M = {"type": "object", "properties": {"metin": S, "dayanak": {"type": "array", "items": S}}, "required": ["metin", "dayanak"], "additionalProperties": False}
SEMA = {
    "type": "object",
    "properties": {"baslik": S, "one_cikanlar": {"type": "array", "items": _M}, "paragraflar": {"type": "array", "items": _M}},
    "required": ["baslik", "one_cikanlar", "paragraflar"],
    "additionalProperties": False,
}
TURLER = ("Kur", "Endeks", "Faiz", "CDS", "Emtia", "Diğer")
BP_TURLERI = {"Faiz", "CDS"}
ARTIS = re.compile(r"\b(yüksel\w*|art(?:tı|ış|arak|mış)\w*|değer kazan\w*|prim yap\w*|tırman\w*|güçlen\w*)", re.I)
AZALIS = re.compile(r"\b(geril\w*|düş(?:tü|üş|erek|müş)\w*|azal\w*|değer kaybet\w*|ger(?:i|ı) çekil\w*|zayıfla\w*)", re.I)
NEDEN = re.compile(r"\b(nedeniyle|sebebiyle|kaynaklı|etkisiyle|dolayı|ardından|sonrasında|sonrası|bağlı olarak|desteğiyle)\b", re.I)
BEKLENTI = re.compile(r"\b(yükselece\w*|düşece\w*|artaca\w*|gerileyece\w*|azalaca\w*|kazanaca\w*|kaybedece\w*|bekliyoruz|öngörüyoruz|tahmin ediyoruz)\b", re.I)
TAVSIYE = re.compile(r"\b(alım fırsatı|satış fırsatı|alım yönlü|satım yönlü|al yönlü|sat yönlü|satın alın\w*|pozisyon alın\w*|öneririz|tavsiye ederiz|"
                     r"hedef fiyat\w*|kaçırmayın|portföye ekle\w*)", re.I)
UYARI = ("Burada yer alan yatırım bilgi, yorum ve değerlendirmeleri yatırım danışmanlığı kapsamında değildir; bilgilendirme amaçlıdır ve kurum "
         "içi kullanım içindir.")


def tur_bul(ad: str, tur: str) -> str:
    t = next((x for x in TURLER if katla(x) == katla(tur)), "")
    if t:
        return t
    k = katla(ad)
    if "cds" in k:
        return "CDS"
    if any(x in k for x in ("faiz", "tahvil", "bono", "gosterge", "repo", "tlref")):
        return "Faiz"
    if "/" in ad or any(x in k for x in ("usd", "eur", "gbp", "dolar", "euro", "sterlin")):
        return "Kur"
    if any(x in k for x in ("bist", "endeks", "nasdaq", "dax", "s p", "dow", "ftse", "nikkei")):
        return "Endeks"
    if any(x in k for x in ("brent", "altin", "gumus", "petrol", "bakir", "ons")):
        return "Emtia"
    return "Diğer"


def _ondalik(v) -> int | None:
    if isinstance(v, str):
        m = re.search(r"[.,](\d+)$", v.strip())
        if m and not re.fullmatch(r"\d{1,3}(?:\.\d{3})+", v.strip()):
            return len(m.group(1))
        return 0
    return None


def tr(v: float, ondalik: int) -> str:
    return f"{abs(v):,.{ondalik}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def isaret(v: float) -> str:
    return "+" if v > 0 else ("−" if v < 0 else "±")


def verileri_oku(yol: Path) -> list[dict]:
    r = kayitlar(yol, {"ad": ("Gösterge", "Enstrüman", "Varlık"), "tur": ("Tür", "Sınıf"), "onceki": ("Önceki Kapanış", "Önceki", "Dünkü Kapanış"),
                       "son": ("Son", "Kapanış", "Son Değer")}, ("ad", "onceki", "son"))
    out = []
    for x in r:
        try:
            o, s = float(sayi_norm(str(x["onceki"]).strip())), float(sayi_norm(str(x["son"]).strip()))
        except ValueError:
            continue
        tur = tur_bul(str(x["ad"]), str(x["tur"]))
        d = _ondalik(x["son"])
        d = d if d is not None else {"Kur": 4, "CDS": 0}.get(tur, 2)
        g = {"id": f"P{len(out) + 1}", "ad": str(x["ad"]).strip(), "tur": tur, "onceki": o, "son": s, "ondalik": d, "degisim": s - o}
        if tur in BP_TURLERI:
            g["bp"] = (s - o) * (100 if tur == "Faiz" else 1)
            g["yuzde"] = None
        else:
            g["bp"] = None
            g["yuzde"] = (s - o) / o * 100 if o else None
        out.append(g)
    return out


def haberleri_oku(yol: Path | None) -> list[str]:
    if not yol:
        return []
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: haber notları bulunamadı.")
    return [s.strip() for s in belge.metin_oku(yol).splitlines() if s.strip()]


def bp_yaz(v: float) -> str:
    return tr(v, 0 if abs(v - round(v)) < 0.05 else 1)


def satir(g: dict) -> str:
    s = f"{g['id']} | {g['ad']} | {g['tur']} | önceki {tr(g['onceki'], g['ondalik'])} | son {tr(g['son'], g['ondalik'])} | "
    if g["bp"] is not None:
        return s + f"değişim {isaret(g['bp'])}{bp_yaz(g['bp'])} baz puan"
    return s + f"değişim {isaret(g['degisim'])}{tr(g['degisim'], g['ondalik'])}" + ("" if g["yuzde"] is None else f" ({isaret(g['yuzde'])}%{tr(g['yuzde'], 2)})")


def rakamlar(gost: list[dict], haber: list[str]) -> dict[str, set[str]]:
    r = {}
    for g in gost:
        x = sayilar(tr(g["onceki"], g["ondalik"])) | sayilar(tr(g["son"], g["ondalik"])) | sayilar(tr(g["degisim"], g["ondalik"])) | sayilar(g["ad"])
        if g["bp"] is not None:
            x |= sayilar(bp_yaz(g["bp"]))
        if g["yuzde"] is not None:
            x |= sayilar(tr(g["yuzde"], 2))
        r[g["id"]] = x
    for i, h in enumerate(haber, 1):
        r[f"H{i}"] = sayilar(h)
    return r


def en_cok(gost: list[dict]) -> list[str]:
    fiyat = [g for g in gost if g["yuzde"] is not None]
    bp = [g for g in gost if g["bp"] is not None]
    out = []
    if fiyat:
        out.append(max(fiyat, key=lambda g: abs(g["yuzde"]))["id"])
    if bp:
        out.append(max(bp, key=lambda g: abs(g["bp"]))["id"])
    return out


def maddeler(t: dict) -> list[tuple[str, dict]]:
    return [("Öne çıkan", m) for m in t["one_cikanlar"]] + [("Paragraf", m) for m in t["paragraflar"]]


def denetle(t: dict, gost: list[dict], haber: list[str], tarih: str) -> list[str]:
    sorun, r = [], rakamlar(gost, haber)
    gd = {g["id"]: g for g in gost}
    kullanilan = set()
    for i, (yer, m) in enumerate(maddeler(t), 1):
        ad = f"M{i} ({yer})"
        if not m["dayanak"]:
            sorun.append(f"{ad}: dayanak (gösterge / haber) yok")
        gec = [d for d in m["dayanak"] if d not in r]
        if gec:
            sorun.append(f"{ad}: olmayan dayanak {gec}")
        kullanilan |= set(m["dayanak"])
        izinli = set().union(set(), *(r.get(d, set()) for d in m["dayanak"])) | sayilar(tarih)
        fazla = sorted(sayilar(m["metin"]) - izinli)
        if fazla:
            sorun.append(f"{ad}: dayandığı gösterge / haberde olmayan sayı: {', '.join(fazla)} (hazır rakamları aynen kullanın)")
        hab = any(d.startswith("H") for d in m["dayanak"])
        for g in (gd[d] for d in m["dayanak"] if d in gd):
            if g["bp"] is None:
                continue
            b = bp_yaz(g["bp"])
            for z in re.finditer(rf"(%\s*{re.escape(b)}\b|\b{re.escape(b)}(?!\s*(?:baz puan|bp|bps|puan)\b)(?![.,]\d))", m["metin"], re.I):
                sorun.append(f"{ad}: {g['ad']} değişimi ({b}) baz puan olarak yazılmalı ('{z.group(0).strip()}')")
                break
        gs = [gd[d] for d in m["dayanak"] if d in gd]
        if len(gs) == 1 and gs[0]["degisim"]:
            art, az = bool(ARTIS.search(m["metin"])), bool(AZALIS.search(m["metin"]))
            if art != az and (art and gs[0]["degisim"] < 0 or az and gs[0]["degisim"] > 0):
                sorun.append(f"{ad}: {gs[0]['ad']} {'geriledi' if gs[0]['degisim'] < 0 else 'yükseldi'}; metindeki yön ters")
        if NEDEN.search(m["metin"]) and not hab:
            sorun.append(f"{ad}: neden-sonuç ifadesi ('{NEDEN.search(m['metin']).group(0)}') bir habere (H…) dayanmıyor")
        if BEKLENTI.search(m["metin"]) and not hab:
            sorun.append(f"{ad}: beklenti / tahmin ifadesi ('{BEKLENTI.search(m['metin']).group(0)}') bir habere dayanmıyor")
        if TAVSIYE.search(m["metin"]):
            sorun.append(f"{ad}: yatırım tavsiyesi sayılabilecek ifade ('{TAVSIYE.search(m['metin']).group(0)}')")
    for d in en_cok(gost):
        if d not in kullanilan:
            sorun.append(f"Genel: günün en çok hareket eden göstergesi ele alınmamış: {d} {gd[d]['ad']}")
    if TAVSIYE.search(t["baslik"]) or sayilar(t["baslik"]) - set().union(set(), *r.values()) - sayilar(tarih):
        sorun.append("Genel: başlıkta tavsiye ifadesi ya da verilerde olmayan sayı var")
    return sorun


def uret(gost, haber, tarih, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    h = "\n".join(f"H{i} | {s}" for i, s in enumerate(haber, 1)) or "(haber notu yok)"
    mesaj = (f"<tarih>{tarih}</tarih>\n<gostergeler>\n" + "\n".join(satir(g) for g in gost) + f"\n</gostergeler>\n<haberler>\n{h}\n</haberler>\n"
             f"<en_cok_hareket>{', '.join(en_cok(gost))}</en_cok_hareket>" + duzeltme_blogu(duzeltme, onceki))
    return llm.json_iste(sistem, mesaj, SEMA)


def rapor_yaz(cikti: Path, t, gost, haber, tarih, kalan, tur) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    baslik(ws, ["Bölüm", "Metin", "Dayanak"], (14, 100, 16))
    ws.append(["Başlık", t["baslik"], ""])
    boya(ws, [2], AI_DOLGU)
    for yer, m in maddeler(t):
        ws.append([yer, m["metin"], ", ".join(m["dayanak"])])
        boya(ws, [2], AI_DOLGU)
    son_not(ws, UYARI, kalan, tur)
    hizala(ws)
    v = wb.create_sheet("Veriler")
    baslik(v, ["No", "Gösterge", "Tür", "Önceki", "Son", "Değişim", "Değişim %", "Değişim (baz puan)"], (5, 26, 9, 12, 12, 12, 11, 16))
    hareket = set(en_cok(gost))
    for g in gost:
        v.append([g["id"], g["ad"], g["tur"], g["onceki"], g["son"], round(g["degisim"], g["ondalik"]), None if g["yuzde"] is None else round(g["yuzde"], 2),
                  None if g["bp"] is None else round(g["bp"], 1)])
        boya(v, [6, 7, 8], YESIL if g["degisim"] > 0 else (KIRMIZI if g["degisim"] < 0 else SARI))
        if g["id"] in hareket:
            boya(v, [2], SARI)
    hizala(v)
    if haber:
        hs = wb.create_sheet("Haberler")
        baslik(hs, ["No", "Haber / Gelişme"], (6, 100))
        for i, s in enumerate(haber, 1):
            hs.append([f"H{i}", s])
        hizala(hs)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    metin = [f"PİYASA GÜNLÜĞÜ — {tarih}", "", t["baslik"], ""] + [f"• {m['metin']}" for m in t["one_cikanlar"]] + [""]
    metin += [m["metin"] + "\n" for m in t["paragraflar"]] + ["Veriler:"] + [satir(g).split(" | ", 1)[1] for g in gost] + ["", UYARI]
    cikti.with_suffix(".txt").write_text("\n".join(metin), encoding="utf-8")
    return "\n".join(metin)


def calistir(girdi: Path, cikti: Path, tarih: str, haber_yolu: Path | None = None, tur: int = 2, evet: bool = False) -> dict:
    gost = verileri_oku(girdi)
    if not gost:
        raise llm.LLMHatasi(f"{girdi.name}: sayısal gösterge yok.")
    haber = haberleri_oku(haber_yolu)
    print(f"[OK] {len(gost)} gösterge · {len(haber)} haber notu · en çok hareket: {', '.join(en_cok(gost))}")
    llm.onay_al(f"{len(gost)} piyasa göstergesi ve {len(haber)} haber notu gönderilecek.", evet)
    t, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(gost, haber, tarih, d, o), lambda x: denetle(x, gost, haber, tarih), tur)
    metin = rapor_yaz(cikti, t, gost, haber, tarih, kalan, yapilan)
    return {"gostergeler": gost, "taslak": t, "kalan": kalan, "tur": yapilan, "metin": metin}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    o = BURASI / "ornek_veri"
    p = argparse.ArgumentParser(description="Gün içi piyasa verilerinden (kur, faiz, CDS, endeks) kısa piyasa özeti hazırlar.")
    p.add_argument("--girdi", type=Path, default=o / "piyasa_verileri.csv", help="Veriler (.xlsx/.csv): Gösterge, Önceki Kapanış, Son; Tür (Kur/Endeks/Faiz/CDS/Emtia)")
    p.add_argument("--haberler", type=Path, default=o / "haberler.txt", help="Günün haber / gelişme notları (.txt/.docx); her satır bir not (isteğe bağlı)")
    p.add_argument("--tarih", default=f"{date.today():%d.%m.%Y}", help="Günlüğün tarihi (varsayılan: bugün)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "piyasa_gunlugu.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.cikti, a.tarih, a.haberler, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {s['taslak']['baslik']}")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()} (+ {a.cikti.with_suffix('.txt').name})")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
