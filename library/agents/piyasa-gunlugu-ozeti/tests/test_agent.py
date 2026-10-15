"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

VERI = KLASOR / "ornek_veri" / "piyasa_verileri.csv"
HABER = KLASOR / "ornek_veri" / "haberler.txt"
TARIH = "14.10.2026"
MESAJLAR = []


def m(metin, *dayanak):
    return {"metin": metin, "dayanak": list(dayanak)}


def taslak(dogru=True) -> dict:
    one = [m("Beklentilerin altında gelen enflasyon verisinin ardından 2 yıllık gösterge tahvil faizi 43 baz puan gerileyerek %39,42'ye indi.", "P5", "H1"),
           m("BIST 100 endeksi %1,64 yükselişle 10.412,85 puandan kapandı.", "P4"),
           m("Brent petrol, OPEC+ üretim artışı kararının ardından %2,18 düşüşle 64,95 dolara geriledi.", "P8", "H3")]
    par = [m("Dolar/TL günü %0,32 artışla 41,9875 seviyesinde tamamlarken EUR/USD paritesi ABD perakende satış verisi sonrası %0,71 geriledi.",
             "P1", "P3", "H2"),
           m("Türkiye 5 yıllık CDS primi 9 baz puan düşüşle 259 baz puana geriledi; 10 yıllık tahvil faizi ise 15 baz puan yükselerek %31,35 oldu.",
             "P7", "P6"),
           m("Ons altın %1,40 yükselişle 4.162,50 dolardan işlem gördü.", "P9")]
    if not dogru:
        one = [m("2 yıllık gösterge tahvil faizi %43 geriledi.", "P5", "H1"), m("BIST 100 endeksi %1,6 düştü.", "P4"),
               m("USD/TRY yükselecek; dolar alım fırsatı sunuyor.", "P1")]
    return {"baslik": "Enflasyon verisi sonrası tahvil faizleri geriledi, BIST 100 yükseldi", "one_cikanlar": one, "paragraflar": par}


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return taslak("<duzeltme>" in kullanici)


class KodTesti(unittest.TestCase):
    def test_hesap(self):
        g = agent.verileri_oku(VERI)
        self.assertEqual([x["tur"] for x in g], ["Kur", "Kur", "Kur", "Endeks", "Faiz", "Faiz", "CDS", "Emtia", "Emtia"])
        self.assertEqual((round(g[4]["bp"]), round(g[6]["bp"]), round(g[3]["yuzde"], 2)), (-43, -9, 1.64))
        self.assertEqual(agent.en_cok(g), ["P8", "P5"])
        self.assertEqual(agent.satir(g[0]), "P1 | USD/TRY | Kur | önceki 41,8520 | son 41,9875 | değişim +0,1355 (+%0,32)")
        self.assertEqual(agent.tur_bul("TR 5Y CDS", ""), "CDS")
        self.assertEqual(agent.tur_bul("Gösterge tahvil 2 yıl", ""), "Faiz")

    def test_denetim(self):
        g, h = agent.verileri_oku(VERI), agent.haberleri_oku(HABER)
        self.assertEqual(agent.denetle(taslak(), g, h, TARIH), [])
        s = agent.denetle(taslak(False), g, h, TARIH)
        self.assertEqual(len(s), 6, s)
        self.assertTrue(any("baz puan" in x and "%43" in x for x in s) and any("1.6" in x for x in s) and any("BIST 100" in x and "ters" in x for x in s)
                        and any("yükselecek" in x for x in s) and any("alım fırsatı" in x for x in s) and any("P8" in x for x in s))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "p.xlsx"
            s = agent.calistir(VERI, cikti, TARIH, HABER, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            self.assertIn("P5 | TR 2Y Gösterge Tahvil | Faiz | önceki 39,85 | son 39,42 | değişim −43 baz puan", MESAJLAR[0])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Veriler", "Haberler"])
            self.assertIn("yatırım danışmanlığı kapsamında değildir", cikti.with_suffix(".txt").read_text(encoding="utf-8"))

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(VERI, Path(tmp) / "p.xlsx", TARIH, HABER, tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 6)
            self.assertEqual(agent.main(["--evet", "--tarih", TARIH, "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--evet", "--girdi", str(VERI.with_name("yok.csv"))]), 1)


if __name__ == "__main__":
    unittest.main()
