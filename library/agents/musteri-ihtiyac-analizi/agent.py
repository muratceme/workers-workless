"""
Müşteri İhtiyaç Analizi — Gerisi Bende AI Agent
Acente Teknik › Sigorta Uzmanı (Acente)

1. Kod görüşme notlarını satır satır numaralar (telefon, e-posta, kimlik no maskelenir) ve acentenin ürün kataloğunu okur
   (Ürün; Branş; Kapsam; Zorunlu).
2. Model müşterinin risklerini (öncelikle, dayandığı satırlarla), katalogdan önerilecek ürünleri (karşıladığı risklerle),
   eksik bilgileri ve müşteriye sorulacak soruları yazar.
3. Kod denetler:
   - Önerilen ürün katalogda var mı; risk ve önerilerin dayanak satırları geçerli mi
   - Her risk bir öneriyle karşılanmış mı (katalogda karşılığı yoksa "katalog dışı" işaretlenmeli)
   - Notlarda araç geçiyorsa zorunlu trafik, kendi konutu geçiyorsa DASK (katalogda "Zorunlu: E" ürünler) ele alınmış mı
   - Notlarda ve katalogda olmayan sayı (prim, teminat tutarı) yazılmış mı
   - "Tam koruma", "her şeyi karşılar", "garanti" gibi aşırı vaat var mı
   Sorun varsa modele geri gönderilir (en fazla --tur kez).

Kullanım:
    python agent.py                                                  # örnek görüşme + örnek katalog
    python agent.py --girdi gorusme.txt --katalog urunler.xlsx
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
K = {"type": "array", "items": {"type": "integer"}}
ONCELIK = {"type": "string", "enum": ["Yüksek", "Orta", "Düşük"]}
SEMA = {
    "type": "object",
    "properties": {
        "musteri_profili": S,
        "riskler": {"type": "array", "items": {"type": "object", "properties": {"risk": S, "oncelik": ONCELIK, "katalog_disi": {"type": "boolean"}, "satirlar": K},
                                               "required": ["risk", "oncelik", "katalog_disi", "satirlar"], "additionalProperties": False}},
        "oneriler": {"type": "array", "items": {"type": "object", "properties": {
            "urun": S, "durum": {"type": "string", "enum": ["Yeni", "Mevcut - yenileme kontrolü", "Mevcut - kapsam gözden geçirme"]},
            "karsiladigi_riskler": K, "gerekce": S, "satirlar": K},
            "required": ["urun", "durum", "karsiladigi_riskler", "gerekce", "satirlar"], "additionalProperties": False}},
        "eksik_bilgiler": {"type": "array", "items": S},
        "sorulacak_sorular": {"type": "array", "items": S},
    },
    "required": ["musteri_profili", "riskler", "oneriler", "eksik_bilgiler", "sorulacak_sorular"],
    "additionalProperties": False,
}
TETIK = {"oto": re.compile(r"\b(arac\w*|otomobil\w*|araba\w*|plaka\w*|kamyonet\w*|motosiklet\w*|minibus\w*|ticari arac\w*)\b"),
         "konut": re.compile(r"\b(kendi (?:evi|evinde|dairesi|dairesinde|konutu)\w*|ev sahibi\w*|evi var|dairesi var|konut sahibi\w*|tapu\w*|mulk sahibi\w*)\b")}
VAAT = re.compile(r"\b(tam koruma|tam guvence|her seyi karsilar\w*|her turlu (?:hasar|zarar|risk)\w*|garanti\w*|hicbir risk\w*|100 koruma|eksiksiz koruma)\b")


def katalogu_oku(yol: Path) -> list[dict]:
    r = kayitlar(yol, {"urun": ("Ürün", "Ürün Adı"), "brans": ("Branş",), "kapsam": ("Kapsam", "Teminatlar", "Açıklama"), "zorunlu": ("Zorunlu",)}, ("urun",))
    return [{"urun": str(x["urun"]).strip(), "brans": str(x["brans"]).strip(), "kapsam": str(x["kapsam"]).strip(),
             "zorunlu": katla(x["zorunlu"]) in {"e", "evet", "x", "1", "zorunlu"}} for x in r if str(x["urun"]).strip()]


def notlari_oku(yol: Path) -> list[str]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: görüşme notu bulunamadı.")
    return [llm.maskele(s.strip()) for s in belge.metin_oku(yol).splitlines() if s.strip()]


def zorunlu_beklenen(l: list[str], katalog: list[dict]) -> list[tuple[dict, int]]:
    """Notlarda tetikleyici geçen zorunlu ürünler ve tetikleyen ilk satır."""
    out = []
    for u in katalog:
        if not u["zorunlu"]:
            continue
        b = katla(u["brans"])
        tetik = next((d for a, d in TETIK.items() if b.startswith(a)), None)
        if not tetik:
            continue
        satir = next((i for i, s in enumerate(l, 1) if tetik.search(katla(s))), None)
        if satir:
            out.append((u, satir))
    return out


def denetle(t: dict, l: list[str], katalog: list[dict]) -> list[str]:
    n, sorun = len(l), []
    adlar = {katla(u["urun"]): u for u in katalog}
    izinli = sayilar("\n".join(l) + " " + " ".join(u["urun"] + " " + u["kapsam"] for u in katalog))

    def satir_kontrol(ad, ks):
        if not ks:
            sorun.append(f"{ad}: dayanak satır yok")
        elif [k for k in ks if not 1 <= k <= n]:
            sorun.append(f"{ad}: notta olmayan satır {[k for k in ks if not 1 <= k <= n]}")

    for i, r in enumerate(t["riskler"], 1):
        satir_kontrol(f"Risk {i}", r["satirlar"])
    karsilanan = set()
    for i, o in enumerate(t["oneriler"], 1):
        if katla(o["urun"]) not in adlar:
            sorun.append(f"Öneri {i}: '{o['urun']}' katalogda yok (yalnız katalogdaki ürün adlarını aynen kullanın)")
        satir_kontrol(f"Öneri {i}", o["satirlar"])
        gec = [k for k in o["karsiladigi_riskler"] if not 1 <= k <= len(t["riskler"])]
        if gec or not o["karsiladigi_riskler"]:
            sorun.append(f"Öneri {i}: karşıladığı risk numarası yok ya da geçersiz {gec}")
        karsilanan |= set(o["karsiladigi_riskler"])
    for i, r in enumerate(t["riskler"], 1):
        if not r["katalog_disi"] and i not in karsilanan:
            sorun.append(f"Risk {i} ('{r['risk'][:40]}'): hiçbir öneriyle karşılanmamış (katalogda ürün yoksa 'katalog_disi' işaretleyin)")
    onerilen = {katla(o["urun"]) for o in t["oneriler"]}
    for u, s in zorunlu_beklenen(l, katalog):
        if katla(u["urun"]) not in onerilen:
            sorun.append(f"Zorunlu: notların {s}. satırına göre '{u['urun']}' ele alınmalı (mevcutsa 'Mevcut - yenileme kontrolü' olarak)")
    metinler = [("Profil", t["musteri_profili"])] + [(f"Risk {i}", r["risk"]) for i, r in enumerate(t["riskler"], 1)]
    metinler += [(f"Öneri {i}", o["gerekce"]) for i, o in enumerate(t["oneriler"], 1)] + [("Soru", s) for s in t["sorulacak_sorular"]]
    for ad, m in metinler:
        fazla = sorted(sayilar(m) - izinli)
        if fazla:
            sorun.append(f"{ad}: notlarda ve katalogda olmayan sayı (prim / teminat uydurulmamalı): {', '.join(fazla)}")
        if VAAT.search(katla(m)):
            sorun.append(f"{ad}: aşırı vaat ('{VAAT.search(katla(m)).group(0)}'); teminat ve istisnalar poliçeye göre belirlenir")
    return sorun


def uret(l: list[str], katalog: list[dict], duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    notlar = "\n".join(f"{i}. {s}" for i, s in enumerate(l, 1))
    kat = "\n".join(f"- {u['urun']} | {u['brans']} | {u['kapsam']}{' | ZORUNLU' if u['zorunlu'] else ''}" for u in katalog)
    mesaj = f"<gorusme_notlari>\n{notlar}\n</gorusme_notlari>\n<urun_katalogu>\n{kat}\n</urun_katalogu>" + duzeltme_blogu(duzeltme, onceki)
    return llm.json_iste(sistem, mesaj, SEMA)


def rapor_yaz(cikti: Path, t, l, katalog, kalan, tur) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Riskler"
    baslik(ws, ["No", "Risk", "Öncelik", "Katalog Dışı", "Karşılayan Ürünler", "Dayanak"], (5, 60, 9, 11, 40, 60))
    renk = {"Yüksek": KIRMIZI, "Orta": SARI, "Düşük": YESIL}
    for i, r in enumerate(t["riskler"], 1):
        urunler = [o["urun"] for o in t["oneriler"] if i in o["karsiladigi_riskler"]]
        ws.append([i, r["risk"], r["oncelik"], "Evet" if r["katalog_disi"] else "", "\n".join(urunler),
                   "\n".join(f"{k}. {l[k - 1]}" for k in r["satirlar"] if 1 <= k <= len(l))])
        boya(ws, [2], AI_DOLGU)
        boya(ws, [3], renk[r["oncelik"]])
    son_not(ws, "Önemli: Ön analizdir. Teklif öncesi teminat, limit, muafiyet ve istisnaları sigorta şirketinin genel ve özel şartlarına göre "
                "müşteriyle birlikte değerlendirin.", kalan, tur)
    hizala(ws)
    o = wb.create_sheet("Öneriler")
    baslik(o, ["Ürün", "Branş", "Durum", "Karşıladığı Riskler", "Gerekçe", "Zorunlu"], (32, 12, 26, 14, 70, 9))
    kat = {katla(u["urun"]): u for u in katalog}
    for x in t["oneriler"]:
        u = kat.get(katla(x["urun"]), {"brans": "?", "zorunlu": False})
        o.append([x["urun"], u["brans"], x["durum"], ", ".join(map(str, x["karsiladigi_riskler"])), x["gerekce"], "Evet" if u["zorunlu"] else ""])
        boya(o, [5], AI_DOLGU)
        if u["zorunlu"]:
            boya(o, [6], SARI)
    hizala(o)
    s = wb.create_sheet("Sorular")
    baslik(s, ["Tür", "Metin"], (16, 100))
    for x in t["eksik_bilgiler"]:
        s.append(["Eksik bilgi", x])
    for x in t["sorulacak_sorular"]:
        s.append(["Sorulacak soru", x])
    hizala(s)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    metin = ["MÜŞTERİ İHTİYAÇ ANALİZİ (ön değerlendirme)", "", t["musteri_profili"], "", "Tespit edilen riskler:"]
    metin += [f"{i}. {r['risk']} ({r['oncelik']})" for i, r in enumerate(t["riskler"], 1)] + ["", "Değerlendirilmesi önerilen ürünler:"]
    metin += [f"- {x['urun']} [{x['durum']}]: {x['gerekce']}" for x in t["oneriler"]] + ["", "Netleştirmemiz gerekenler:"]
    metin += [f"- {x}" for x in t["sorulacak_sorular"]]
    metin += ["", "Bu değerlendirme bilgilendirme amaçlıdır; teminat kapsamı, limitler ve istisnalar poliçe genel ve özel şartlarına göre belirlenir."]
    cikti.with_suffix(".txt").write_text("\n".join(metin), encoding="utf-8")
    return "\n".join(metin)


def calistir(girdi: Path, katalog_yolu: Path, cikti: Path, tur: int = 2, evet: bool = False) -> dict:
    l = notlari_oku(girdi)
    katalog = katalogu_oku(katalog_yolu)
    if not l or not katalog:
        raise llm.LLMHatasi("Görüşme notu ya da ürün kataloğu boş.")
    z = zorunlu_beklenen(l, katalog)
    print(f"[OK] {len(l)} not satırı · {len(katalog)} katalog ürünü · ele alınması gereken zorunlu: {', '.join(u['urun'] for u, _ in z) or '-'}")
    llm.onay_al(f"{len(l)} satırlık görüşme notu (telefon / e-posta / kimlik no maskeli) ve ürün kataloğu gönderilecek.", evet)
    t, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(l, katalog, d, o), lambda x: denetle(x, l, katalog), tur)
    metin = rapor_yaz(cikti, t, l, katalog, kalan, yapilan)
    return {"satirlar": l, "katalog": katalog, "taslak": t, "kalan": kalan, "tur": yapilan, "metin": metin}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    o = BURASI / "ornek_veri"
    p = argparse.ArgumentParser(description="Görüşme notlarından müşterinin risklerini ve katalogdan önerilecek sigorta ürünlerini içeren ihtiyaç analizi hazırlar.")
    p.add_argument("--girdi", type=Path, default=o / "gorusme_notu.txt", help="Görüşme notu (.txt/.docx); her satır bir bilgi")
    p.add_argument("--katalog", type=Path, default=o / "urun_katalogu.csv", help="Acentenin ürün kataloğu (.xlsx/.csv): Ürün; Branş, Kapsam, Zorunlu (E/H)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "ihtiyac_analizi.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.katalog, a.cikti, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(s['taslak']['riskler'])} risk · {len(s['taslak']['oneriler'])} ürün önerisi")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()} (+ {a.cikti.with_suffix('.txt').name})")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
