"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

NOT = KLASOR / "ornek_veri" / "gorusme_notu.txt"
KATALOG = KLASOR / "ornek_veri" / "urun_katalogu.csv"
MESAJLAR = []


def r(risk, oncelik, *s, disi=False):
    return {"risk": risk, "oncelik": oncelik, "katalog_disi": disi, "satirlar": list(s)}


def o(urun, durum, riskler, gerekce, *s):
    return {"urun": urun, "durum": durum, "karsiladigi_riskler": riskler, "gerekce": gerekce, "satirlar": list(s)}


def analiz(dogru=True) -> dict:
    t = {"musteri_profili": "Kendi konutunda oturan, aracı ve küçük bir kafesi olan aile; aylık taksitli ödemeyi tercih ediyor.",
         "riskler": [r("Konut binasının deprem hasarı; DASK durumu bilinmiyor", "Yüksek", 1, 2), r("Ev eşyasının yangın, hırsızlık ve su baskını riski", "Orta", 1, 4),
                     r("Aracın kendi hasarı ve çalınma riski (kasko yok)", "Yüksek", 3), r("Aracın üçüncü kişilere verebileceği zararlar", "Düşük", 3),
                     r("Kafedeki yaklaşık 400.000 TL değerindeki ekipmanın yangın ve hırsızlık riski", "Yüksek", 6),
                     r("3 çalışanın iş kazası nedeniyle işverene yöneltebileceği tazminat talepleri", "Orta", 5),
                     r("Özel hastane fark ücretleri", "Orta", 7), r("Kiralık dükkânda kira sözleşmesindeki sigorta yükümlülükleri belirsiz", "Düşük", 6, disi=True)],
         "oneriler": [o("DASK (Zorunlu Deprem Sigortası)", "Yeni", [1], "Kendi konutu için DASK poliçesi bulunmuyorsa yaptırılması zorunludur.", 1, 2),
                      o("Konut Sigortası", "Yeni", [2], "Ev eşyası ve değerli eşyalar için yangın ve hırsızlık teminatı.", 1, 4),
                      o("Kasko", "Yeni", [3], "2021 model aracın kendi hasarları için.", 3),
                      o("Zorunlu Trafik Sigortası", "Mevcut - yenileme kontrolü", [4], "Mevcut poliçenin bitiş tarihi takip edilmeli.", 3),
                      o("İşyeri Sigortası", "Yeni", [5], "Yaklaşık 400.000 TL değerindeki ekipman için.", 6),
                      o("İşveren Mali Sorumluluk Sigortası", "Yeni", [6], "Kafedeki 3 çalışan için iş kazası tazminat taleplerine karşı.", 5),
                      o("Tamamlayıcı Sağlık Sigortası", "Yeni", [7], "Anlaşmalı özel hastane tercihi nedeniyle fark ücretleri için.", 7)],
         "eksik_bilgiler": ["Konutun tapu durumu ve mevcut DASK poliçesi"], "sorulacak_sorular": ["Kafenin metrekaresi ve faaliyet konusu nedir?"]}
    if not dogru:
        t["oneriler"] = [x for x in t["oneriler"] if not x["urun"].startswith("DASK")]
        t["oneriler"][1]["urun"] = "Kasko Plus"
        t["oneriler"][3]["gerekce"] += " Aylık 850 TL primle tam koruma sağlar."
    return t


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return analiz("<duzeltme>" in kullanici)


class KodTesti(unittest.TestCase):
    def test_zorunlu(self):
        l, k = agent.notlari_oku(NOT), agent.katalogu_oku(KATALOG)
        self.assertEqual(len(k), 8)
        self.assertEqual(k[3]["kapsam"], "Bina ve ev eşyası; yangın, hırsızlık, su baskını (poliçe şartlarına göre)")
        self.assertEqual([(u["urun"], s) for u, s in agent.zorunlu_beklenen(l, k)], [("Zorunlu Trafik Sigortası", 3), ("DASK (Zorunlu Deprem Sigortası)", 1)])
        self.assertEqual(agent.zorunlu_beklenen(["Şirket aracı kullanıyor."], k)[0][0]["urun"], "Zorunlu Trafik Sigortası")
        self.assertEqual(agent.zorunlu_beklenen(["Kiracı olarak oturuyor."], k), [])

    def test_denetim(self):
        l, k = agent.notlari_oku(NOT), agent.katalogu_oku(KATALOG)
        self.assertEqual(agent.denetle(analiz(), l, k), [])
        s = agent.denetle(analiz(False), l, k)
        self.assertEqual(len(s), 5, s)
        self.assertTrue(any("Kasko Plus" in x for x in s) and any(x.startswith("Risk 1 ") for x in s) and any("DASK" in x and "Zorunlu" in x for x in s)
                        and any("850" in x for x in s) and any("tam koruma" in x for x in s))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "i.xlsx"
            s = agent.calistir(NOT, KATALOG, cikti, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            self.assertIn("DASK (Zorunlu Deprem Sigortası) | Konut | Konut binasının deprem ve deprem kaynaklı hasarı (limitler dahilinde) | ZORUNLU", MESAJLAR[0])
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Riskler", "Öneriler", "Sorular"])
            self.assertEqual(wb["Riskler"]["E2"].value, "DASK (Zorunlu Deprem Sigortası)")
            self.assertIn("poliçe genel ve özel şartlarına", cikti.with_suffix(".txt").read_text(encoding="utf-8"))

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(NOT, KATALOG, Path(tmp) / "i.xlsx", tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 5)
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--evet", "--katalog", str(KATALOG.with_name("yok.csv"))]), 1)


if __name__ == "__main__":
    unittest.main()
