"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

VERI = KLASOR / "ornek_veri"
GUVENILIR = ["ornekfirma.com.tr", "ornekbank.com.tr"]
MESAJLAR = []


def g(gosterge, aciklama, alinti=""):
    return {"gosterge": gosterge, "aciklama": aciklama, "alinti": alinti}


def analiz(kimlik, dogru=True) -> dict:
    if kimlik == "E1":
        a = {"risk": "Yüksek", "sinif": "Oltalama (kimlik bilgisi)", "ozet": "Banka taklidi yapan, şifre toplamaya yönelik oltalama e-postası.",
             "gerekceler": [g("G2", "DMARC başarısız"), g("G5", "Gönderen alan adı bankayı taklit ediyor"), g("G6", "Yanıt adresi başka bir sahte alan adı"),
                            g("G7", "Görünen ad banka adını kullanıyor"), g("G8", "Bağlantı metni gerçek adresi gizliyor"), g("G9", "Bağlantı IP adresine gidiyor"),
                            g("G10", "Çift uzantılı HTML eki"), g("G11", "Aciliyet baskısı", "Hesabınız 24 saat içinde askıya alınacaktır")],
             "kullaniciya_tavsiye": "Bağlantıya tıklamayın, eki açmayın ve hiçbir bilgi girmeyin. E-postayı BT'ye bildirin.",
             "bt_aksiyonlari": ["ornekbank-guvenlik[.]com ve mail-ornekbank[.]xyz alan adlarını engelleyin", "198[.]51[.]100[.]23 adresini engelleyin"]}
        if not dogru:
            a["risk"] = "Orta"
            a["gerekceler"] = a["gerekceler"][:1] + a["gerekceler"][4:5]
            a["bt_aksiyonlari"][0] = "mail-ornekbank.xyz alan adını engelleyin"
            a["kullaniciya_tavsiye"] = "Emin değilseniz bağlantıya tıklayın ve adres çubuğunu kontrol edin."
        return a
    if kimlik == "E2":
        return {"risk": "Düşük", "sinif": "Meşru", "ozet": "Kurum içi eğitim duyurusu; kimlik doğrulama başarılı.", "gerekceler": [],
                "kullaniciya_tavsiye": "Duyuru kurum içinden geliyor; portala her zamanki yer iminizden ulaşabilirsiniz.", "bt_aksiyonlari": []}
    a = {"risk": "Yüksek", "sinif": "CEO / ödeme dolandırıcılığı", "ozet": "Yönetici adına kişisel e-posta adresinden gizli ve acil havale talebi.",
         "gerekceler": [g("G2", "Ödeme talebi", "bugün mesai bitmeden 185.000 TL havale yapılması gerekiyor"), g("G3", "Ücretsiz e-posta servisi"),
                        g("", "Gizlilik baskısı", "Konuyu kimseyle paylaşma")],
         "kullaniciya_tavsiye": "Ödeme yapmayın; talebi yöneticinin kayıtlı telefon numarasından arayarak teyit edin.",
         "bt_aksiyonlari": ["Gönderen adresini engelleyin", "Finans ekibini benzer taleplere karşı uyarın"]}
    if not dogru:
        a["gerekceler"][0]["alinti"] = "Hemen 200.000 TL gönder"
    return a


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    kimlik = re.search(r'<eposta id="(E\d+)"', kullanici).group(1)
    return analiz(kimlik, "<duzeltme>" in kullanici)


def epostalar():
    out = []
    for i, d in enumerate(agent.dosyalar(VERI), 1):
        e = agent.eml_oku(d)
        e["kimlik"], e["gostergeler"] = f"E{i}", agent.gostergeler(e, GUVENILIR)
        e["kod_riski"] = agent.kod_riski(e["gostergeler"])
        out.append(e)
    return out


class KodTesti(unittest.TestCase):
    def test_gostergeler(self):
        e = epostalar()
        self.assertEqual([x["kod_riski"] for x in e], ["Yüksek", "Düşük", "Orta"])
        yuksek = {x["tur"] for x in e[0]["gostergeler"] if x["seviye"] == "Yüksek"}
        self.assertEqual(yuksek, {"Kimlik doğrulama", "Benzer alan adı", "Görünen ad", "Bağlantı", "Ek"})
        self.assertEqual(len(e[0]["linkler"]), 1)
        self.assertEqual(e[1]["gostergeler"], [])
        self.assertTrue(any("gmail.com" in x["aciklama"] for x in e[2]["gostergeler"]))
        self.assertEqual(agent.kayitli_alan("www.mail.ornekbank.com.tr"), "ornekbank.com.tr")
        self.assertEqual(agent.etkisiz("http://198.51.100.23/a"), "hxxp://198[.]51[.]100[.]23/a")

    def test_benzer_alan(self):
        e = {"gonderen": "a@0rnekbank.com", "gorunen_ad": "", "reply_to": "", "return_path": "", "auth": "spf=pass", "konu": "", "govde": "",
             "linkler": [("https://xn--rnekbank-xxa.com/", ""), ("https://bit.ly/abc", "")], "ekler": [{"ad": "fatura.zip"}, {"ad": "kurulum.exe"}]}
        s = [(x["seviye"], x["tur"]) for x in agent.gostergeler(e, GUVENILIR)]
        self.assertEqual(s, [("Yüksek", "Benzer alan adı"), ("Yüksek", "Benzer alan adı"), ("Orta", "Bağlantı"), ("Orta", "Ek"), ("Yüksek", "Ek")])

    def test_denetim(self):
        e = epostalar()
        dogru = {x["kimlik"]: analiz(x["kimlik"]) for x in e}
        self.assertEqual(agent.denetle(e, dogru, GUVENILIR), [])
        s = agent.denetle(e, {x["kimlik"]: analiz(x["kimlik"], False) for x in e}, GUVENILIR)
        self.assertEqual(len(s), 5, s)
        self.assertTrue(any("E1" in x and "altında" in x for x in s) and any("E1" in x and "G5" in x for x in s)
                        and any("mail-ornekbank.xyz" in x and "etkisiz" in x for x in s) and any("tiklayin" in x for x in s)
                        and any(x.startswith("E3") and "alıntı" in x for x in s))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "o.xlsx"
            s = agent.calistir(VERI, cikti, GUVENILIR, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (5, [], 1))
            self.assertEqual(sorted(re.search(r'id="(E\d)"', m).group(1) for m in MESAJLAR[3:]), ["E1", "E3"])
            self.assertNotIn("muhasebe@ornekfirma", MESAJLAR[0])
            self.assertNotIn("http://198", MESAJLAR[0])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Göstergeler", "Bağlantılar"])

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(VERI / "1_hesap_dogrulama.eml", Path(tmp) / "o.xlsx", GUVENILIR, tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 4)
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--evet", "--girdi", str(VERI / "yok.txt")]), 1)


if __name__ == "__main__":
    unittest.main()
