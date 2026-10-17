"""
Oltalama E-postası Analizi — Gerisi Bende AI Agent
Bilgi Teknolojileri › Bilgi Güvenliği Uzmanı

1. Kod .eml dosyalarını açar (bağlantılara gitmez, ekleri çalıştırmaz) ve göstergeleri çıkarır:
   - SPF / DKIM / DMARC sonucu (Authentication-Results)
   - Reply-To ve Return-Path alan adı gönderenden farklı mı
   - Görünen ad kurum / marka adı taşıyıp başka alan adından mı geliyor
   - Güvenilir alan adına benzeyen (marka adı içeren, harf oynanmış, punycode) alan adı
   - Bağlantı metninde görünen adres ile gerçek adres farklı mı; IP adresine ya da kısaltıcıya bağlantı
   - Riskli ek (çalıştırılabilir, betik, ISO / LNK / HTML, makrolu Office, çift uzantı)
   - Aciliyet / parola / ödeme talebi ifadeleri; ücretsiz e-posta servisinden ödeme talebi
   Kod göstergelere göre en düşük risk seviyesini belirler.
2. Model her e-posta için risk seviyesi, sınıf, gerekçe (gösterge numarası ve metinden alıntıyla), kullanıcıya tavsiye
   ve BT aksiyonlarını yazar.
3. Kod denetler:
   - Model riski, göstergelerin gerektirdiği seviyenin altında mı
   - Yüksek seviyeli her gösterge gerekçede kullanılmış mı; gösterge numaraları ve alıntılar geçerli mi
   - Şüpheli bağlantı / alan adı etkisizleştirilmeden (hxxp, [.]) yazılmış mı
   - Tavsiyede "bağlantıya tıklayın", "eki açın" gibi yönlendirme var mı
   Sorun varsa yalnız o e-postanın analizi modele geri gönderilir (en fazla --tur kez).

Kullanım:
    python agent.py                                                 # örnek .eml dosyaları
    python agent.py --girdi supheli.eml --guvenilir "firmam.com.tr,bankam.com.tr"
    python agent.py --girdi eml_klasoru/
"""
from __future__ import annotations

import argparse
import csv
import difflib
import json
import re
import sys
from datetime import date, datetime
from email import policy
from email.parser import BytesParser
from email.utils import parseaddr
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
SEVIYE = {"Düşük": 0, "Orta": 1, "Yüksek": 2}
SEMA = {
    "type": "object",
    "properties": {
        "risk": {"type": "string", "enum": list(SEVIYE)},
        "sinif": {"type": "string", "enum": ["Oltalama (kimlik bilgisi)", "Zararlı ek / yazılım", "CEO / ödeme dolandırıcılığı", "Spam / reklam", "Meşru", "Belirsiz"]},
        "ozet": S,
        "gerekceler": {"type": "array", "items": {"type": "object", "properties": {"gosterge": S, "aciklama": S, "alinti": S},
                                                  "required": ["gosterge", "aciklama", "alinti"], "additionalProperties": False}},
        "kullaniciya_tavsiye": S,
        "bt_aksiyonlari": {"type": "array", "items": S},
    },
    "required": ["risk", "sinif", "ozet", "gerekceler", "kullaniciya_tavsiye", "bt_aksiyonlari"],
    "additionalProperties": False,
}
TR_IKINCI = {"com", "net", "org", "gov", "edu", "gen", "biz", "info", "av", "bel", "k12", "web", "tv", "tsk", "pol", "bbs", "name", "co", "ac"}
KISALTICI = {"bit.ly", "tinyurl.com", "t.co", "goo.gl", "is.gd", "cutt.ly", "ow.ly", "rebrand.ly", "shorturl.at", "t.ly", "rb.gy"}
UCRETSIZ = {"gmail.com", "hotmail.com", "outlook.com", "yahoo.com", "yandex.com", "yandex.com.tr", "icloud.com", "mail.ru", "proton.me", "protonmail.com", "hotmail.com.tr"}
ZARARLI_UZANTI = {".exe", ".scr", ".js", ".jse", ".vbs", ".vbe", ".bat", ".cmd", ".ps1", ".iso", ".img", ".lnk", ".hta", ".html", ".htm", ".msi", ".jar", ".wsf", ".com", ".pif"}
ORTA_UZANTI = {".docm", ".xlsm", ".pptm", ".xlam", ".zip", ".rar", ".7z", ".gz"}
ACILIYET = re.compile(r"\b(acil\w*|hemen|derhal|24 saat|48 saat|askiya al\w*|kapatilacak\w*|bloke|son uyari|dogrula\w*|urgent|immediately|suspend\w*|verify)\b")
KIMLIK = re.compile(r"\b(sifre\w*|parola\w*|kullanici adi\w*|kart bilgi\w*|cvv|tek kullanimlik|sms kod\w*|password|login|giris yap\w*)\b")
ODEME = re.compile(r"\b(havale|eft|iban|odeme\w*|transfer|hesap numara\w*|wire|payment|hediye kart\w*)\b")
YONLENDIRME = re.compile(r"\b(tiklayin\w*|tiklayiniz|eki acin\w*|ekini acin\w*|eki acabilirsiniz|baglantiya girin\w*|baglantiya girerek|sifrenizi girin\w*|bilgilerinizi girin\w*|"
                         r"baglantidan giris yap\w*)\b")
URL = re.compile(r"\b(?:https?://|www\.)[^\s<>\"')\]]+", re.I)
ALAN_DESEN = re.compile(r"\b(?:[a-z0-9-]+\.)+[a-z]{2,}\b", re.I)


class _Html(HTMLParser):
    def __init__(self):
        super().__init__()
        self.linkler, self.metin, self._a = [], [], None

    def handle_starttag(self, tag, attrs):
        if tag == "a":
            self._a = [dict(attrs).get("href", ""), ""]
        if tag in {"br", "p", "div", "tr", "li"}:
            self.metin.append("\n")

    def handle_endtag(self, tag):
        if tag == "a" and self._a:
            self.linkler.append((self._a[0], self._a[1].strip()))
            self._a = None

    def handle_data(self, data):
        self.metin.append(data)
        if self._a:
            self._a[1] += data


def alan(adres_ya_da_url: str) -> str:
    s = str(adres_ya_da_url or "").strip().lower()
    if "@" in s and "://" not in s:
        s = s.rsplit("@", 1)[1]
    if "://" not in s:
        s = "http://" + s
    try:
        h = urlsplit(s).hostname or ""
    except ValueError:
        h = ""
    return h.strip(".")


def kayitli_alan(h: str) -> str:
    """'www.mail.ornek.com.tr' → 'ornek.com.tr'; 'a.ornek.com' → 'ornek.com'."""
    p = h.split(".")
    if len(p) >= 3 and len(p[-1]) == 2 and p[-2] in TR_IKINCI:
        return ".".join(p[-3:])
    return ".".join(p[-2:])


def etiket(h: str) -> str:
    return kayitli_alan(h).split(".")[0]


def ip_mi(h: str) -> bool:
    return bool(re.fullmatch(r"\d{1,3}(?:\.\d{1,3}){3}", h))


def etkisiz(s: str) -> str:
    """Bağlantıyı tıklanamaz yapar: http → hxxp, nokta → [.]"""
    return re.sub(r"^http", "hxxp", s, flags=re.I).replace(".", "[.]")


def eml_oku(yol: Path) -> dict:
    with yol.open("rb") as f:
        m = BytesParser(policy=policy.default).parse(f)
    duz, html, ekler = [], [], []
    for p in m.walk():
        if p.is_multipart():
            continue
        ad = p.get_filename()
        if ad or p.get_content_disposition() == "attachment":
            ekler.append({"ad": ad or "(adsız)", "tur": p.get_content_type()})
            continue
        try:
            icerik = p.get_content()
        except (LookupError, ValueError):
            icerik = p.get_payload(decode=True).decode("utf-8", "replace") if p.get_payload(decode=True) else ""
        if p.get_content_type() == "text/plain":
            duz.append(str(icerik))
        elif p.get_content_type() == "text/html":
            html.append(str(icerik))
    linkler, govde = [], "\n".join(duz)
    for h in html:
        ph = _Html()
        ph.feed(h)
        linkler += ph.linkler
        if not duz:
            govde += re.sub(r"\n\s*\n+", "\n", "".join(ph.metin))
    bilinen = {x[0] for x in linkler} | {x[1] for x in linkler}  # HTML bağlantısının görünen metni ayrıca sayılmaz
    linkler += [(u, "") for u in dict.fromkeys(URL.findall(govde)) if u not in bilinen]
    ad, adres = parseaddr(str(m.get("From", "")))
    return {"dosya": yol.name, "konu": str(m.get("Subject", "")), "tarih": str(m.get("Date", "")), "gorunen_ad": ad, "gonderen": adres,
            "reply_to": parseaddr(str(m.get("Reply-To", "")))[1], "return_path": parseaddr(str(m.get("Return-Path", "")))[1],
            "auth": " ".join(str(x) for x in m.get_all("Authentication-Results", [])), "govde": govde.strip(), "linkler": linkler, "ekler": ekler}


def gostergeler(e: dict, guvenilir: list[str]) -> list[dict]:
    g = []

    def ekle(seviye, tur, aciklama):
        g.append({"id": f"G{len(g) + 1}", "seviye": seviye, "tur": tur, "aciklama": aciklama})

    gon = alan(e["gonderen"])
    guv_etiket = {etiket(x): x for x in guvenilir}
    guvenilir_mi = lambda h: any(h == x or h.endswith("." + x) for x in guvenilir)  # noqa: E731

    a = e["auth"].lower()
    for mek, seviye in (("spf", "Orta"), ("dkim", "Orta"), ("dmarc", "Yüksek")):
        m = re.search(rf"\b{mek}=(\w+)", a)
        if m and m.group(1) in {"fail", "softfail", "permerror"}:
            ekle(seviye, "Kimlik doğrulama", f"{mek.upper()} sonucu: {m.group(1)}")
    if not a:
        ekle("Düşük", "Kimlik doğrulama", "Authentication-Results başlığı yok; SPF / DKIM / DMARC sonucu bilinmiyor")
    for baslik_, adres, seviye in (("Reply-To", e["reply_to"], "Orta"), ("Return-Path", e["return_path"], "Düşük")):
        if adres and gon and kayitli_alan(alan(adres)) != kayitli_alan(gon):
            ekle(seviye, "Başlık uyumsuzluğu", f"{baslik_} alan adı ({etkisiz(alan(adres))}) gönderenden ({etkisiz(gon)}) farklı")

    def benzer(h: str, yer: str):
        if not h or guvenilir_mi(h) or ip_mi(h):
            return
        if "xn--" in h:
            ekle("Yüksek", "Benzer alan adı", f"{yer}: punycode alan adı ({etkisiz(h)}); harfleri benzer görünen sahte adres olabilir")
            return
        et = etiket(h)
        for ge, tam in guv_etiket.items():
            if ge in et or (len(et) > 3 and difflib.SequenceMatcher(None, et, ge).ratio() >= 0.8):
                ekle("Yüksek", "Benzer alan adı", f"{yer}: {etkisiz(h)} güvenilir {tam} adresine benziyor")
                return

    benzer(gon, "Gönderen")
    if e["reply_to"]:
        benzer(alan(e["reply_to"]), "Reply-To")
    gad = re.sub(r"[^a-z0-9]", "", katla(e["gorunen_ad"]))
    if gon and not guvenilir_mi(gon) and any(ge and ge.replace("-", "") in gad for ge in guv_etiket):
        ekle("Yüksek", "Görünen ad", f"Görünen ad ('{e['gorunen_ad']}') güvenilir bir kurumu çağrıştırıyor ama adres {etkisiz(gon)}")
    gorulen = set()
    for href, yazi in e["linkler"]:
        h = alan(href)
        if not h or (href, yazi) in gorulen:
            continue
        gorulen.add((href, yazi))
        gorunen = ALAN_DESEN.search(yazi or "")
        if gorunen and kayitli_alan(alan(gorunen.group(0))) != kayitli_alan(h):
            ekle("Yüksek", "Bağlantı", f"Bağlantı metni {etkisiz(gorunen.group(0))} gösteriyor ama gerçek adres {etkisiz(href)}")
        if ip_mi(h):
            ekle("Yüksek", "Bağlantı", f"Bağlantı doğrudan IP adresine gidiyor: {etkisiz(href)}")
        elif h in KISALTICI:
            ekle("Orta", "Bağlantı", f"Kısaltılmış bağlantı (gerçek hedef gizli): {etkisiz(href)}")
        else:
            benzer(h, "Bağlantı")
    for x in e["ekler"]:
        ad = x["ad"].lower()
        uz = Path(ad).suffix
        parcalar = ad.split(".")
        if len(parcalar) >= 3 and "." + parcalar[-2] in {".pdf", ".doc", ".docx", ".xls", ".xlsx", ".jpg", ".png", ".txt"} and uz in ZARARLI_UZANTI | ORTA_UZANTI:
            ekle("Yüksek", "Ek", f"Çift uzantılı ek: {x['ad']}")
        elif uz in ZARARLI_UZANTI:
            ekle("Yüksek", "Ek", f"Çalıştırılabilir / betik / kısayol türünde ek: {x['ad']}")
        elif uz in ORTA_UZANTI:
            ekle("Orta", "Ek", f"Makro içerebilen Office dosyası ya da arşiv: {x['ad']}")
    metin = katla(e["konu"] + " " + e["govde"])
    for desen, ad in ((ACILIYET, "Aciliyet / tehdit"), (KIMLIK, "Kimlik bilgisi talebi"), (ODEME, "Ödeme talebi")):
        k = sorted(set(desen.findall(metin)))
        if k:
            ekle("Orta" if ad != "Aciliyet / tehdit" else "Düşük", "İçerik", f"{ad} ifadeleri: {', '.join(k[:6])}")
    if kayitli_alan(gon) in UCRETSIZ and ODEME.search(metin):
        ekle("Orta", "İçerik", f"Ücretsiz e-posta servisinden ({gon}) ödeme talebi; kurum içi yönetici adına dolandırıcılık olabilir")
    return g


def kod_riski(g: list[dict]) -> str:
    if any(x["seviye"] == "Yüksek" for x in g):
        return "Yüksek"
    if any(x["seviye"] == "Orta" for x in g):
        return "Orta"
    return "Düşük"


def supheli_alanlar(e: dict, guvenilir: list[str]) -> set[str]:
    hs = {alan(h) for h, _ in e["linkler"]} | {alan(e["gonderen"]), alan(e["reply_to"]), alan(e["return_path"])}
    return {h for h in hs if h and not any(h == x or h.endswith("." + x) for x in guvenilir)}


def denetle(epostalar: list[dict], sonuc: dict, guvenilir: list[str]) -> list[str]:
    sorun = []
    for e in epostalar:
        ad, a = e["kimlik"], sonuc.get(e["kimlik"])
        if not a:
            sorun.append(f"{ad}: analiz yazılmadı")
            continue
        ids = {x["id"] for x in e["gostergeler"]}
        if SEVIYE[a["risk"]] < SEVIYE[e["kod_riski"]]:
            sorun.append(f"{ad}: risk '{a['risk']}' göstergelerin gerektirdiği '{e['kod_riski']}' seviyesinin altında")
        if a["sinif"] == "Meşru" and e["kod_riski"] == "Yüksek":
            sorun.append(f"{ad}: yüksek seviyeli gösterge varken 'Meşru' denmiş")
        kullanilan = {x["gosterge"] for x in a["gerekceler"]}
        for x in a["gerekceler"]:
            if x["gosterge"] and x["gosterge"] not in ids:
                sorun.append(f"{ad}: gerekçede olmayan gösterge {x['gosterge']}")
            if x["alinti"] and not metinde_var(x["alinti"], e["konu"] + "\n" + e["govde"]):
                sorun.append(f"{ad}: alıntı e-postada yok: '{x['alinti'][:50]}'")
        eksik = [x["id"] for x in e["gostergeler"] if x["seviye"] == "Yüksek" and x["id"] not in kullanilan]
        if eksik:
            sorun.append(f"{ad}: yüksek seviyeli gösterge gerekçede kullanılmamış: {', '.join(eksik)}")
        cikti = " ".join([a["ozet"], a["kullaniciya_tavsiye"]] + a["bt_aksiyonlari"] + [x["aciklama"] for x in a["gerekceler"]])
        acik = sorted({h for h in supheli_alanlar(e, guvenilir) if h.lower() in cikti.lower()} | set(re.findall(r"https?://\S+", cikti)))
        if acik:
            sorun.append(f"{ad}: bağlantı / alan adı etkisizleştirilmeden yazılmış: {', '.join(acik)} (hxxp ve [.] kullanın)")
        if YONLENDIRME.search(katla(a["kullaniciya_tavsiye"])):
            sorun.append(f"{ad}: tavsiyede bağlantıya / eke yönlendirme var ('{YONLENDIRME.search(katla(a['kullaniciya_tavsiye'])).group(0)}')")
    return sorun


def ozet_metni(e: dict) -> str:
    govde = llm.maskele(e["govde"])[:6000]
    link = "\n".join(f"- metin: '{y[:80]}' → adres: {etkisiz(h)}" for h, y in e["linkler"][:30]) or "(yok)"
    ek = ", ".join(x["ad"] for x in e["ekler"]) or "(yok)"
    g = "\n".join(f"{x['id']} [{x['seviye']}] {x['tur']}: {x['aciklama']}" for x in e["gostergeler"]) or "(kod gösterge bulmadı)"
    return (f'<eposta id="{e["kimlik"]}">\nKonu: {e["konu"]}\nGörünen ad: {e["gorunen_ad"]}\nGönderen alan adı: {etkisiz(alan(e["gonderen"]))}\n'
            f"Reply-To alan adı: {etkisiz(alan(e['reply_to'])) or '-'}\nEkler: {ek}\nBağlantılar:\n{link}\n<govde>\n{govde}\n</govde>\n"
            f"<gostergeler en_dusuk_risk=\"{e['kod_riski']}\">\n{g}\n</gostergeler>\n</eposta>")


def uret(epostalar: list[dict], guvenilir: list[str], duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    sonuc, hedef = {}, epostalar
    if duzeltme is not None:  # düzeltmede yalnız sorunlu e-postalar yeniden analiz edilir
        sorunlu = {s.split(":")[0] for s in duzeltme}
        sonuc.update({a: b for a, b in (onceki or {}).items() if a not in sorunlu})
        hedef = [e for e in epostalar if e["kimlik"] in sorunlu]
    for e in hedef:
        d = [s for s in (duzeltme or []) if s.split(":")[0] == e["kimlik"]]
        mesaj = (f"<guvenilir_alanlar>{', '.join(guvenilir) or '(verilmedi)'}</guvenilir_alanlar>\n" + ozet_metni(e)
                 + duzeltme_blogu(d, (onceki or {}).get(e["kimlik"])))
        sonuc[e["kimlik"]] = llm.json_iste(sistem, mesaj, SEMA)
    return sonuc


def dosyalar(girdi: Path) -> list[Path]:
    if girdi.is_dir():
        return sorted(girdi.glob("*.eml"))
    if not girdi.exists():
        raise llm.LLMHatasi(f"{girdi}: dosya ya da klasör bulunamadı.")
    if girdi.suffix.lower() != ".eml":
        raise llm.LLMHatasi(f"{girdi.name}: .eml dosyası bekleniyor (Outlook'ta iletiyi 'Farklı Kaydet' ya da sürükle-bırak ile .eml / .msg → .eml).")
    return [girdi]


def rapor_yaz(cikti: Path, epostalar, sonuc, kalan, tur) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    baslik(ws, ["No", "Dosya", "Konu", "Görünen Ad", "Gönderen Alan Adı", "Kod Riski", "AI Riski", "Sınıf", "Özet", "Kullanıcıya Tavsiye", "BT Aksiyonları"],
           (5, 20, 36, 22, 26, 10, 10, 22, 60, 50, 50))
    renk = {"Yüksek": KIRMIZI, "Orta": SARI, "Düşük": YESIL}
    for e in epostalar:
        a = sonuc.get(e["kimlik"])
        satir = [e["kimlik"], e["dosya"], e["konu"], e["gorunen_ad"], etkisiz(alan(e["gonderen"])), e["kod_riski"]]
        ws.append(satir + ([a["risk"], a["sinif"], a["ozet"], a["kullaniciya_tavsiye"], "\n".join(a["bt_aksiyonlari"])] if a else ["", "", "AI yanıt vermedi"]))
        boya(ws, [6], renk[e["kod_riski"]])
        if a:
            boya(ws, [8, 9, 10, 11], AI_DOLGU)
            boya(ws, [7], renk[a["risk"]])
    son_not(ws, "Önemli: Bağlantılar ziyaret edilmedi, ekler açılmadı. Şüpheli iletiyi silmeden önce kurumunuzun olay müdahale "
                "prosedürüne göre saklayın ve bildirin.", kalan, tur)
    hizala(ws)
    g = wb.create_sheet("Göstergeler")
    baslik(g, ["E-posta", "Gösterge", "Seviye", "Tür", "Açıklama", "AI Gerekçesi"], (8, 9, 9, 18, 70, 60))
    for e in epostalar:
        a = sonuc.get(e["kimlik"]) or {"gerekceler": []}
        for x in e["gostergeler"]:
            ws_ger = "\n".join(y["aciklama"] for y in a["gerekceler"] if y["gosterge"] == x["id"])
            g.append([e["kimlik"], x["id"], x["seviye"], x["tur"], x["aciklama"], ws_ger])
            boya(g, [3], renk[x["seviye"]])
    hizala(g)
    b = wb.create_sheet("Bağlantılar")
    baslik(b, ["E-posta", "Görünen Metin", "Gerçek Adres (etkisiz)", "Alan Adı"], (8, 40, 70, 30))
    for e in epostalar:
        for h, y in e["linkler"]:
            b.append([e["kimlik"], y, etkisiz(h), etkisiz(alan(h))])
    hizala(b)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(girdi: Path, cikti: Path, guvenilir: list[str], tur: int = 2, evet: bool = False) -> dict:
    epostalar = []
    for i, d in enumerate(dosyalar(girdi), 1):
        e = eml_oku(d)
        e["kimlik"] = f"E{i}"
        e["gostergeler"] = gostergeler(e, guvenilir)
        e["kod_riski"] = kod_riski(e["gostergeler"])
        epostalar.append(e)
    if not epostalar:
        raise llm.LLMHatasi(f"{girdi}: .eml dosyası yok.")
    for e in epostalar:
        print(f"[OK] {e['kimlik']} {e['dosya']}: {len(e['gostergeler'])} gösterge · kod riski {e['kod_riski']}")
    llm.onay_al(f"{len(epostalar)} e-postanın konu, gövde (e-posta / telefon / IBAN maskeli), bağlantı ve gösterge özeti gönderilecek.", evet)
    sonuc, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(epostalar, guvenilir, d, o), lambda x: denetle(epostalar, x, guvenilir), tur)
    rapor_yaz(cikti, epostalar, sonuc, kalan, yapilan)
    return {"epostalar": epostalar, "sonuc": sonuc, "kalan": kalan, "tur": yapilan}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    p = argparse.ArgumentParser(description="Şüpheli e-postaların başlık, bağlantı, ek ve içerik göstergelerini inceleyip risk değerlendirmesi yazar.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri", help=".eml dosyası ya da .eml dosyalarının bulunduğu klasör")
    p.add_argument("--guvenilir", default="ornekfirma.com.tr,ornekbank.com.tr",
                   help="Kurumunuzun ve taklit edilmesi olası kurumların alan adları (virgülle)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "oltalama_analizi.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    guvenilir = [x.strip().lower() for x in a.guvenilir.split(",") if x.strip()]
    try:
        s = calistir(a.girdi, a.cikti, guvenilir, a.tur, a.evet)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    for e in s["epostalar"]:
        r = s["sonuc"].get(e["kimlik"])
        print(f"[{'!' if r and r['risk'] != 'Düşük' else 'OK'}] {e['kimlik']} {e['konu'][:50]}: {r['risk'] + ' · ' + r['sinif'] if r else 'analiz yok'}")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
