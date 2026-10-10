"""
Sosyal Medya İçerik Takvimi — Workers / Workless AI Agent
Pazarlama › Dijital Pazarlama Uzmanı

Kampanya ve özel günlere göre aylık sosyal medya içerik takvimi ve gönderi metni taslakları hazırlar:
1. Kod takvimi kurar: platform başına haftalık gönderi sayısına göre günleri seçer; özel günleri (sabit tarihli günler,
   Anneler / Babalar Günü, Kasım indirim cuması; dinî bayramlar dosyadan), kampanya dönemlerini ve içerik sütunlarını
   (eğitici, ürün, sosyal kanıt...) sırayla dağıtır. Anma ve farkındalık günleri "satış dili yok" olarak işaretlenir.
2. Model her gönderi için metin, hashtag, görsel fikri ve eylem çağrısı taslağı yazar (marka tonu, platform sınırı ve
   kampanya mesajına göre).
3. Kod taslakları denetler: platform karakter sınırı, hashtag sayısı, yasaklı ifadeler, kanıt gerektiren iddialar
   ("en iyi", "garanti", "%100"...), anma / farkındalık gününde satış ifadesi, kampanya mesajında olmayan oran veya
   fiyat. Sorunlu gönderiler işaretlenir; yayın kararı insandadır.
--model-yok ile yalnız takvim iskeleti çıkar (API gerekmez).

Kullanım:
    python agent.py                                                  # örnek: Kasım 2026, 3 platform
    python agent.py --ay 2026-11 --profil marka_profili.txt --kampanyalar kampanyalar.xlsx --ozel-gunler ozel_gunler.csv
    python agent.py --ay 2026-12 --profil marka_profili.txt --model-yok
"""
from __future__ import annotations

import argparse
import calendar
import csv
import re
import sys
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import llm

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
PAKET = 20
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
GUNLER = ["Pazartesi", "Salı", "Çarşamba", "Perşembe", "Cuma", "Cumartesi", "Pazar"]
AYLAR = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]
# platform → (karakter sınırı, önerilen en çok hashtag, kesin hashtag sınırı)
PLATFORMLAR = {"instagram": ("Instagram", 2200, 8, 30), "x": ("X", 280, 2, 10), "twitter": ("X", 280, 2, 10), "linkedin": ("LinkedIn", 3000, 5, 30),
               "facebook": ("Facebook", 2000, 3, 30), "tiktok": ("TikTok", 2200, 5, 30)}
GUN_DAGILIMI = {1: [2], 2: [1, 3], 3: [0, 2, 4], 4: [0, 1, 3, 4], 5: [0, 1, 2, 3, 4], 6: [0, 1, 2, 3, 4, 5], 7: [0, 1, 2, 3, 4, 5, 6]}
SUTUNLAR = ["Eğitici / ipucu", "Ürün / hizmet tanıtımı", "Sosyal kanıt (müşteri yorumu, referans)", "Perde arkası / ekip", "Etkileşim (soru, anket)"]
# (ay, gün, ad, tür) — tür: kutlama | anma | farkindalik
SABIT_GUNLER = [
    (1, 1, "Yılbaşı", "kutlama"), (2, 14, "Sevgililer Günü", "kutlama"), (3, 8, "Dünya Kadınlar Günü", "farkindalik"), (3, 14, "Tıp Bayramı", "kutlama"),
    (3, 18, "Çanakkale Zaferi ve Şehitleri Anma Günü", "anma"), (3, 22, "Dünya Su Günü", "farkindalik"),
    (4, 23, "Ulusal Egemenlik ve Çocuk Bayramı", "kutlama"), (5, 1, "Emek ve Dayanışma Günü", "kutlama"),
    (5, 19, "Atatürk'ü Anma, Gençlik ve Spor Bayramı", "kutlama"), (6, 5, "Dünya Çevre Günü", "farkindalik"),
    (7, 15, "Demokrasi ve Millî Birlik Günü", "anma"), (8, 30, "Zafer Bayramı", "kutlama"), (10, 4, "Hayvanları Koruma Günü", "farkindalik"),
    (10, 29, "Cumhuriyet Bayramı", "kutlama"), (11, 10, "Atatürk'ü Anma Günü", "anma"), (11, 20, "Dünya Çocuk Hakları Günü", "farkindalik"),
    (11, 24, "Öğretmenler Günü", "kutlama"), (11, 25, "Kadına Yönelik Şiddete Karşı Uluslararası Mücadele Günü", "farkindalik"),
    (12, 3, "Dünya Engelliler Günü", "farkindalik"), (12, 31, "Yılbaşı gecesi", "kutlama"),
]
TUR_ADI = {"kutlama": "Kutlama", "anma": "Anma", "farkindalik": "Farkındalık"}
IDDIA = ["en iyi", "en ucuz", "en kaliteli", "en hizli", "1 numara", "bir numara", "lider", "rakipsiz", "garanti", "garantili", "kesin sonuc", "mucize", "yuzde yuz",
         "risksiz", "tedavi", "iyilestirir", "zayiflatir", "sifir risk"]
SATIS = ["indirim", "kampanya", "firsat", "satin al", "siparis", "kupon", "hemen al", "stoklarla sinirli", "ucretsiz kargo", "sepet"]


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").translate(_TR).lower()).split())


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


# ----------------------------------------------------------------------------
# Profil, kampanyalar, özel günler
# ----------------------------------------------------------------------------

def profil_oku(yol: Path) -> dict:
    p = {"satirlar": [], "platformlar": {}, "sutunlar": [], "yasakli": [], "hashtagler": []}
    for kod in ("utf-8-sig", "cp1254"):
        try:
            ham = yol.read_text(encoding=kod)
            break
        except UnicodeDecodeError:
            continue
    for satir in ham.splitlines():
        if not satir.strip():
            continue
        ad, _, deger = satir.partition(":")
        k, parcalar = katla(ad), [x.strip() for x in re.split(r"[;,]", deger) if x.strip()]
        if k.startswith("platform"):
            for x in parcalar:
                a, _, n = x.partition("=")
                if katla(a) not in PLATFORMLAR:
                    raise ValueError(f"{yol.name}: platform '{a.strip()}' tanınmıyor ({', '.join(sorted({v[0] for v in PLATFORMLAR.values()}))})")
                p["platformlar"][PLATFORMLAR[katla(a)][0]] = max(1, min(7, int(n))) if n.strip().isdigit() else 3
        elif k.startswith("icerik sutun"):
            p["sutunlar"] = parcalar
        elif k.startswith("yasakli"):
            p["yasakli"] = parcalar
        elif "hashtag" in k:
            p["hashtagler"] = [x if x.startswith("#") else "#" + x for x in parcalar]
        p["satirlar"].append(satir.strip())
    if not p["platformlar"]:
        raise ValueError(f"{yol.name}: 'Platformlar: Instagram=3; LinkedIn=2' biçiminde bir satır gerekli")
    p["sutunlar"] = p["sutunlar"] or SUTUNLAR
    return p


def kampanyalari_oku(yol: Path | None) -> list[dict]:
    if not yol:
        return []
    satirlar = tablo_oku(yol)
    b = [katla(c) for c in satirlar[0]]
    lst = []
    for r in satirlar[1:]:
        al = lambda *adlar: next((r[b.index(a)] for a in adlar if a in b and b.index(a) < len(r)), None)  # noqa: E731
        bas, bit = tarih(al("baslangic", "baslangic tarihi")), tarih(al("bitis", "bitis tarihi"))
        if not bas or not metin(al("kampanya")):
            continue
        pl = {PLATFORMLAR[katla(x)][0] for x in re.split(r"[,/|]", metin(al("platformlar", "platform"))) if katla(x) in PLATFORMLAR}
        lst.append({"bas": bas, "bit": bit or bas, "ad": metin(al("kampanya")), "mesaj": metin(al("mesaj", "aciklama")), "urun": metin(al("urun", "urun hizmet")),
                    "platformlar": pl})
    return lst


def n_inci_gun(yil: int, ay: int, haftanin_gunu: int, n: int) -> date:
    ilk = date(yil, ay, 1)
    return ilk + timedelta(days=(haftanin_gunu - ilk.weekday()) % 7 + 7 * (n - 1))


def ozel_gunler(yil: int, ay: int, dosya: Path | None = None) -> dict[date, tuple[str, str]]:
    g = {date(yil, m, d): (ad, tur) for m, d, ad, tur in SABIT_GUNLER if m == ay}
    if ay == 5:
        g[n_inci_gun(yil, 5, 6, 2)] = ("Anneler Günü", "kutlama")              # Mayıs'ın 2. Pazarı
    if ay == 6:
        g[n_inci_gun(yil, 6, 6, 3)] = ("Babalar Günü", "kutlama")              # Haziran'ın 3. Pazarı
    if ay == 11:
        g[n_inci_gun(yil, 11, 3, 4) + timedelta(days=1)] = ("Kasım indirim cuması", "kutlama")   # 4. Perşembe'den sonraki Cuma
    if dosya:
        for r in tablo_oku(dosya)[1:]:
            t = tarih(r[0]) if r else None
            if t and (t.year, t.month) == (yil, ay) and len(r) > 1 and metin(r[1]):
                tur = katla(r[2]) if len(r) > 2 else ""
                g[t] = (metin(r[1]), "anma" if tur.startswith("anma") else "farkindalik" if tur.startswith("fark") else "kutlama")
    return dict(sorted(g.items()))


# ----------------------------------------------------------------------------
# Takvim (kod)
# ----------------------------------------------------------------------------

@dataclass
class Gonderi:
    id: str
    tarih: date
    platform: str
    tur: str                 # Özel gün | Kampanya | İçerik
    konu: str
    brif: str = ""
    satis_yok: bool = False
    kampanya: str = ""       # rakam denetiminde esas alınacak kampanya
    metin: str = ""
    hashtagler: list[str] = field(default_factory=list)
    gorsel: str = ""
    cta: str = ""
    kontrol: list[tuple[str, str]] = field(default_factory=list)

    @property
    def tam_metin(self) -> str:
        return (self.metin + ("\n\n" + " ".join(self.hashtagler) if self.hashtagler else "")).strip()


def takvim_kur(yil: int, ay: int, profil: dict, kampanyalar: list[dict], gunler: dict) -> list[Gonderi]:
    son = calendar.monthrange(yil, ay)[1]
    ana = next(iter(profil["platformlar"]))
    plan = []
    for g in range(1, son + 1):
        d = date(yil, ay, g)
        for pl, n in profil["platformlar"].items():
            if d.weekday() in GUN_DAGILIMI[n]:
                plan.append((d, pl))
        if d in gunler and not any(x[0] == d for x in plan):
            plan.append((d, ana))                         # özel gün, planlı gönderi olmayan güne denk geldiyse ana platformda paylaşılır
    plan.sort(key=lambda x: (x[0], list(profil["platformlar"]).index(x[1])))
    gonderiler, sutun_i, kampanya_sayac = [], 0, {}
    for i, (d, pl) in enumerate(plan, 1):
        aktif = [k for k in kampanyalar if k["bas"] <= d <= k["bit"] and (not k["platformlar"] or pl in k["platformlar"])]
        if d in gunler:
            ad, tur = gunler[d]
            brif = f"{TUR_ADI[tur]} günü: {ad}."
            if tur != "kutlama":
                brif += " Satış, indirim veya ürün tanıtımı YOK; saygılı ve sade."
            elif aktif:
                brif += f" Süren kampanya: {aktif[0]['ad']}. Mesaj: {aktif[0]['mesaj']}."
            gonderiler.append(Gonderi(f"G{i:03d}", d, pl, "Özel gün", ad, brif, tur != "kutlama", aktif[0]["ad"] if aktif and tur == "kutlama" else ""))
            continue
        if aktif:
            k = aktif[0]
            anahtar = (k["ad"], pl)
            kampanya_sayac[anahtar] = kampanya_sayac.get(anahtar, 0) + 1
            if kampanya_sayac[anahtar] % 2 == 1:           # kampanya döneminde her iki gönderiden biri kampanyaya ayrılır
                kalan = (k["bit"] - d).days
                gonderiler.append(Gonderi(f"G{i:03d}", d, pl, "Kampanya", k["ad"],
                                          f"Kampanya: {k['ad']}. Mesaj: {k['mesaj']}." + (f" Ürün: {k['urun']}." if k["urun"] else "")
                                          + (" Kampanyanın son günü." if kalan == 0 else f" Bitişe {kalan} gün." if kalan <= 3 else ""), kampanya=k["ad"]))
                continue
        sutun = profil["sutunlar"][sutun_i % len(profil["sutunlar"])]
        sutun_i += 1
        gonderiler.append(Gonderi(f"G{i:03d}", d, pl, "İçerik", sutun, f"İçerik sütunu: {sutun}."))
    return gonderiler


# ----------------------------------------------------------------------------
# Model ve denetim
# ----------------------------------------------------------------------------

SEMA = {"type": "object", "properties": {"gonderiler": {"type": "array", "items": {
    "type": "object",
    "properties": {"id": {"type": "string"}, "metin": {"type": "string"}, "hashtagler": {"type": "array", "items": {"type": "string"}},
                   "gorsel_fikri": {"type": "string"}, "cta": {"type": "string"}},
    "required": ["id", "metin", "hashtagler", "gorsel_fikri", "cta"], "additionalProperties": False}}},
        "required": ["gonderiler"], "additionalProperties": False}


def platform_bilgisi(ad: str) -> tuple[str, int, int, int]:
    return PLATFORMLAR[katla(ad)]


def taslaklari_yaz(gonderiler: list[Gonderi], profil: dict) -> None:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    for i in range(0, len(gonderiler), PAKET):
        parca = gonderiler[i:i + PAKET]
        satirlar = ["<marka_profili>", *profil["satirlar"], "</marka_profili>", "<gonderiler>"]
        for g in parca:
            _, sinir, onerilen, _ = platform_bilgisi(g.platform)
            satirlar.append(f'<gonderi id="{g.id}" tarih="{g.tarih:%d.%m.%Y} {GUNLER[g.tarih.weekday()]}" platform="{g.platform}" karakter_siniri="{sinir}" '
                            f'en_cok_hashtag="{onerilen}" tur="{g.tur}" satis_dili="{"yok" if g.satis_yok else "serbest"}">{g.brif}</gonderi>')
        satirlar.append("</gonderiler>")
        yanit = llm.json_iste(sistem, "\n".join(satirlar), SEMA)
        harita = {g.id: g for g in parca}
        for s in yanit.get("gonderiler", []):
            g = harita.get(s.get("id"))
            if g is None or g.metin:
                continue
            g.metin = metin(s.get("metin"))
            g.hashtagler = list(dict.fromkeys("#" + re.sub(r"[^\wçğıöşüÇĞİÖŞÜ]", "", h) for h in s.get("hashtagler", []) if re.sub(r"[^\wçğıöşüÇĞİÖŞÜ]", "", h)))
            g.gorsel, g.cta = metin(s.get("gorsel_fikri")), metin(s.get("cta"))


def rakamlar(s: str) -> set[str]:
    """Metindeki yüzde ve TL tutarları: '%20', '20%', '1.500 TL' → {'%20', '1500tl'}."""
    k = s.replace(".", "").replace(" ", "").lower().replace("₺", "tl")
    bulunan = {"%" + m for m in re.findall(r"%(\d+)", k)} | {"%" + m for m in re.findall(r"(\d+)%", k)}
    return bulunan | {m + "tl" for m in re.findall(r"(\d+)(?:,\d+)?tl", k)}


def denetle(g: Gonderi, profil: dict, kampanya_metni: str) -> None:
    if not g.metin:
        g.kontrol.append(("Yüksek", "Model bu gönderi için taslak döndürmedi"))
        return
    ad, sinir, onerilen, kesin = platform_bilgisi(g.platform)
    uzunluk = len(g.tam_metin)
    if uzunluk > sinir:
        g.kontrol.append(("Yüksek", f"{ad} karakter sınırı aşıldı: {uzunluk} / {sinir} (hashtagler dahil)"))
    if len(g.hashtagler) > kesin:
        g.kontrol.append(("Yüksek", f"Hashtag sayısı {len(g.hashtagler)}; {ad} sınırı {kesin}"))
    elif len(g.hashtagler) > onerilen:
        g.kontrol.append(("Bilgi", f"Hashtag sayısı {len(g.hashtagler)}; {ad} için önerilen en çok {onerilen}"))
    k = " " + katla(g.metin + " " + g.cta) + " "
    for y in profil["yasakli"]:
        if katla(y) and " " + katla(y) in k:
            g.kontrol.append(("Yüksek", f"Yasaklı ifade: '{y}'"))
    bulunan = [i for i in IDDIA if " " + i in k] + (["%100"] if "%100" in k.replace(" ", "") else [])
    if bulunan:
        g.kontrol.append(("Orta", "Kanıt gerektiren iddia: " + ", ".join(f"'{i}'" for i in bulunan) + ". Belgelenemiyorsa çıkarın"))
    if g.satis_yok:
        satis = [w for w in SATIS if " " + w in k]
        if satis or rakamlar(g.metin):
            g.kontrol.append(("Yüksek", "Anma / farkındalık gününde satış ifadesi: " + ", ".join(satis + sorted(rakamlar(g.metin)))))
    else:
        izinli = rakamlar(kampanya_metni + " " + " ".join(profil["satirlar"]))
        fazla = rakamlar(g.metin + " " + g.cta) - izinli
        if fazla:
            g.kontrol.append(("Orta", "Kampanya mesajında veya profilde olmayan oran / tutar: " + ", ".join(sorted(fazla))))


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
TUR_RENK = {"Özel gün": "E8F0FE", "Kampanya": "FFF4CE", "İçerik": "FFFFFF"}
AI = PatternFill("solid", fgColor="EDE7FF")
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, basliklar, genislik=()):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def rapor_yaz(cikti: Path, s: dict) -> None:
    wb = Workbook()
    tk = wb.active
    tk.title = "Takvim"
    _baslik(tk, ["No", "Tarih", "Gün", "Platform", "Tür", "Konu", "Brif", "Metin (taslak)", "Hashtagler", "Görsel Fikri", "Eylem Çağrısı", "Karakter", "Sınır",
                 "Kontrol", "Onay (Evet / Düzelt)"], (6, 11, 10, 10, 10, 26, 40, 70, 30, 40, 24, 8, 7, 50, 14))
    for g in s["gonderiler"]:
        sinir = platform_bilgisi(g.platform)[1]
        tk.append([g.id, g.tarih, GUNLER[g.tarih.weekday()], g.platform, g.tur, g.konu, g.brif, g.metin, " ".join(g.hashtagler), g.gorsel, g.cta,
                   len(g.tam_metin) if g.metin else None, sinir, "\n".join(f"[{o}] {a}" for o, a in g.kontrol), ""])
        r = tk.max_row
        tk.cell(r, 2).number_format = "DD.MM.YYYY"
        tk.cell(r, 5).fill = PatternFill("solid", fgColor=TUR_RENK[g.tur])
        for j in (7, 8, 9, 10, 11, 14):
            tk.cell(r, j).alignment = UST
        if g.metin:
            for j in (8, 9, 10, 11):
                tk.cell(r, j).fill = AI
        if g.kontrol:
            enk = "Yüksek" if any(o == "Yüksek" for o, _ in g.kontrol) else "Orta" if any(o == "Orta" for o, _ in g.kontrol) else "Bilgi"
            tk.cell(r, 14).fill = PatternFill("solid", fgColor=RENK[enk])
        tk.cell(r, 15).fill = PatternFill("solid", fgColor="FFF4CE")
    tk.auto_filter.ref = f"A1:O{tk.max_row}"

    oz = wb.create_sheet("Özet")
    _baslik(oz, ["Gösterge", "Değer"], (40, 50))
    g_ = s["gonderiler"]
    for a, v in [("Ay", f"{AYLAR[s['ay'] - 1]} {s['yil']}"), ("Gönderi sayısı", len(g_))] + \
                [(f"  {pl}", sum(x.platform == pl for x in g_)) for pl in s["profil"]["platformlar"]] + \
                [("Özel gün gönderisi", sum(x.tur == "Özel gün" for x in g_)), ("Kampanya gönderisi", sum(x.tur == "Kampanya" for x in g_)),
                 ("İçerik gönderisi", sum(x.tur == "İçerik" for x in g_)), ("Kontrolde sorun çıkan gönderi", sum(any(o != "Bilgi" for o, _ in x.kontrol) for x in g_)),
                 ("Model", s["model"])]:
        oz.append([a, v])
    oz.append([])
    oz.append(["Mor hücreler yapay zekâ taslağıdır. Yayından önce her gönderiyi okuyun; marka, hukuk ve güncel olaylar açısından uygunluğu sizin kararınızdır."])

    og = wb.create_sheet("Özel Günler")
    _baslik(og, ["Tarih", "Gün", "Özel Gün", "Tür", "Not"], (11, 10, 50, 12, 50))
    for d, (ad, tur) in s["gunler"].items():
        og.append([d, GUNLER[d.weekday()], ad, TUR_ADI[tur], "Satış dili kullanılmaz" if tur != "kutlama" else ""])
        og.cell(og.max_row, 1).number_format = "DD.MM.YYYY"

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Gönderi", "Tarih", "Platform", "Açıklama"], (9, 8, 11, 10, 110))
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    for o, g, a in sorted(((o, g, a) for g in g_ for o, a in g.kontrol), key=lambda x: (sira[x[0]], x[1].id)):
        uy.append([o, g.id, g.tarih, g.platform, a])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[o])
        uy.cell(uy.max_row, 3).number_format = "DD.MM.YYYY"
    for u in s["uyarilar"]:
        uy.append([u["onem"], "", None, "", u["aciklama"]])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def csv_yaz(yol: Path, gonderiler: list[Gonderi]) -> None:
    with open(yol, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Tarih", "Platform", "Tür", "Konu", "Metin", "Hashtagler", "Görsel Fikri", "Eylem Çağrısı"])
        for g in gonderiler:
            w.writerow([g.tarih.strftime("%d.%m.%Y"), g.platform, g.tur, g.konu, g.metin, " ".join(g.hashtagler), g.gorsel, g.cta])


def calistir(yil: int, ay: int, profil_yolu: Path, cikti: Path, *, kampanya_yolu: Path | None = None, ozel_gun_yolu: Path | None = None, model: bool = True,
             evet: bool = False) -> dict:
    profil = profil_oku(profil_yolu)
    kampanyalar = kampanyalari_oku(kampanya_yolu)
    gunler = ozel_gunler(yil, ay, ozel_gun_yolu)
    gonderiler = takvim_kur(yil, ay, profil, kampanyalar, gunler)
    uy = []
    ay_ici = [k for k in kampanyalar if k["bas"] <= date(yil, ay, calendar.monthrange(yil, ay)[1]) and k["bit"] >= date(yil, ay, 1)]
    for k in ay_ici:
        if not any(g.tur == "Kampanya" and g.konu == k["ad"] for g in gonderiler):
            uy.append({"onem": "Orta", "aciklama": f"'{k['ad']}' kampanyası için bu ay planlı gönderi günü yok (dönem çok kısa veya platform eşleşmiyor)"})
    if not ozel_gun_yolu:
        uy.append({"onem": "Bilgi", "aciklama": "Dinî bayramlar ve sektöre özel günler takvimde yok; --ozel-gunler dosyasıyla ekleyin"})
    print(f"[OK] {AYLAR[ay - 1]} {yil}: {len(gonderiler)} gönderi · {len(gunler)} özel gün · {len(ay_ici)} kampanya")
    durum = "kullanılmadı (--model-yok)"
    if model:
        llm.onay_al(f"Marka profili, {len(ay_ici)} kampanya mesajı ve {len(gonderiler)} gönderi brifi metin taslağı için gönderilecek. "
                    "Profil ve kampanya dosyalarında gizli bilgi olmadığından emin olun.", evet)
        taslaklari_yaz(gonderiler, profil)
        kampanya_metni = {k["ad"]: f"{k['ad']} {k['mesaj']} {k['urun']}" for k in kampanyalar}
        for g in gonderiler:
            denetle(g, profil, kampanya_metni.get(g.kampanya, ""))
        durum = llm.kullanim_ozeti()
    s = {"yil": yil, "ay": ay, "profil": profil, "kampanyalar": kampanyalar, "gunler": gunler, "gonderiler": gonderiler, "uyarilar": uy, "model": durum}
    rapor_yaz(cikti, s)
    s["csv"] = cikti.with_suffix(".csv")
    csv_yaz(s["csv"], gonderiler)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Kampanya ve özel günlere göre aylık sosyal medya içerik takvimi ve gönderi taslakları hazırlar.")
    p.add_argument("--ay", help="Takvim ayı YYYY-AA (örnek veride 2026-11)")
    p.add_argument("--profil", type=Path, default=ORNEK / "marka_profili.txt",
                   help="Marka profili: Marka, Sektör, Hedef Kitle, Ton, Platformlar (Instagram=3; LinkedIn=2), İçerik Sütunları, Yasaklı İfadeler, Sabit Hashtagler")
    p.add_argument("--kampanyalar", type=Path, help="Başlangıç, Bitiş, Kampanya, Mesaj, Ürün, Platformlar")
    p.add_argument("--ozel-gunler", type=Path, help="Ek özel günler: Tarih, Ad, Tür (Kutlama / Anma / Farkındalık) — dinî bayramlar, sektör günleri")
    p.add_argument("--model-yok", action="store_true", help="Modele gönderme; yalnız takvim iskeleti")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "icerik_takvimi.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    ornek = a.profil == ORNEK / "marka_profili.txt"
    kamp = a.kampanyalar or (ORNEK / "kampanyalar.csv" if ornek else None)
    ozel = a.ozel_gunler or (ORNEK / "ozel_gunler.csv" if ornek else None)
    for y in (a.profil, kamp, ozel):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    m = re.fullmatch(r"(\d{4})-(\d{1,2})", a.ay or ("2026-11" if ornek else ""))
    if not m or not 1 <= int(m[2]) <= 12:
        print("[X] --ay YYYY-AA biçiminde olmalı (ör. 2026-11)")
        return 2
    try:
        s = calistir(int(m[1]), int(m[2]), a.profil, a.cikti, kampanya_yolu=kamp, ozel_gun_yolu=ozel, model=not a.model_yok, evet=a.evet)
    except (llm.LLMHatasi, ValueError) as h:
        print(f"[X] {h}")
        return 1
    sorunlu = [g for g in s["gonderiler"] if any(o == "Yüksek" for o, _ in g.kontrol)]
    for g in sorunlu:
        print(f"[!] {g.id} {g.tarih:%d.%m} {g.platform}: " + "; ".join(a_ for o, a_ in g.kontrol if o == "Yüksek"))
    print(f"[OK] Takvim: {a.cikti.resolve()}")
    print(f"[OK] CSV: {s['csv'].resolve()}")
    if not a.model_yok:
        print(f"[i] Kullanım: {s['model']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
