"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import re
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

PROFIL, KAMPANYA, OZEL = agent.ORNEK / "marka_profili.txt", agent.ORNEK / "kampanyalar.csv", agent.ORNEK / "ozel_gunler.csv"
MESAJLAR = []
OZEL_METIN = {
    "G001": ("Şehrin en iyi kahvesi bizde.", ["NitelikliKahve"]),                            # kanıt gerektiren iddia
    "G002": ("x" * 300, []),                                                                  # X sınırı 280
    "G003": ("Kaliteli kahve artık çok ucuz!", ["kahve"]),                                    # yasaklı ifade
    "G004": ("Atölyemizden bir kare.", ["a", "b", "c", "d", "e", "f", "g"]),                  # LinkedIn için çok hashtag
    "G007": ("Kış harmanımız satışta. İlk hafta 250 g paketlerde %15 indirim.", ["#KışHarmanı"]),
    "G008": ("Saygıyla anıyoruz. Bugüne özel indirim için sitemize bekleriz.", []),          # anma gününde satış
    "G011": ("Kış harmanında %25 indirim sizi bekliyor.", []),                                # kampanyada %15 yazıyor
    "G024": ("Demleme ekipmanlarında %20 indirim sürüyor.", ["#Kasım"]),
}


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append((sistem, kullanici))
    kayitlar = []
    for i in re.findall(r'<gonderi id="([^"]+)"', kullanici):
        if i == "G005":
            continue                                                                          # model bir gönderiyi atladı
        m, h = OZEL_METIN.get(i, ("Bugün demleme önerimiz: suyu kaynadıktan sonra bir dakika bekletin.", ["kahve", "#kahve", "demleme ipucu"]))
        kayitlar.append({"id": i, "metin": m, "hashtagler": h, "gorsel_fikri": "Ürün fotoğrafı", "cta": "Yorumlarda bize yazın"})
    return {"gonderiler": kayitlar + [{"id": "G999", "metin": "x", "hashtagler": [], "gorsel_fikri": "", "cta": ""}]}


class TakvimTesti(unittest.TestCase):
    def test_ozel_gunler(self):
        self.assertEqual(agent.n_inci_gun(2026, 5, 6, 2), date(2026, 5, 10))                  # Anneler Günü 2026
        self.assertEqual(agent.n_inci_gun(2026, 6, 6, 3), date(2026, 6, 21))                  # Babalar Günü 2026
        g = agent.ozel_gunler(2026, 11, OZEL)
        self.assertEqual(g[date(2026, 11, 27)][0], "Kasım indirim cuması")                    # 4. Perşembe (26) + 1
        self.assertEqual(g[date(2026, 11, 10)], ("Atatürk'ü Anma Günü", "anma"))
        self.assertEqual(g[date(2026, 11, 11)][0], "11.11 alışveriş günü")
        self.assertIn(date(2026, 5, 10), agent.ozel_gunler(2026, 5))

    def test_iskelet(self):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(2026, 11, PROFIL, Path(tmp) / "t.xlsx", kampanya_yolu=KAMPANYA, ozel_gun_yolu=OZEL, model=False)
            g = s["gonderiler"]
            self.assertEqual(MESAJLAR, [])
            self.assertEqual(len(g), 25)
            self.assertEqual({x.platform: sum(y.platform == x.platform for y in g) for x in g}, {"Instagram": 13, "X": 8, "LinkedIn": 4})
            anma = next(x for x in g if x.tarih == date(2026, 11, 10))
            self.assertTrue(anma.satis_yok and anma.tur == "Özel gün")
            kampanya = [x for x in g if x.tur == "Kampanya"]
            self.assertEqual({x.konu for x in kampanya}, {"Kış harmanı lansmanı", "Kasım indirim haftası"})
            self.assertTrue(all(x.platform != "LinkedIn" for x in kampanya if x.konu == "Kış harmanı lansmanı"))   # yalnız Instagram ve X
            cuma = next(x for x in g if x.tarih == date(2026, 11, 27))
            self.assertEqual(cuma.kampanya, "Kasım indirim haftası")
            self.assertTrue(s["csv"].read_text(encoding="utf-8-sig").startswith("Tarih;Platform"))

    def test_rakamlar(self):
        self.assertEqual(agent.rakamlar("%20 indirim, 1.500 TL üzeri; ikinci üründe 50% ve 99,90 TL"), {"%20", "%50", "1500tl", "99tl"})


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_uctan_uca(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            s = agent.calistir(2026, 11, PROFIL, cikti, kampanya_yolu=KAMPANYA, ozel_gun_yolu=OZEL, evet=True)
            self.assertEqual(len(MESAJLAR), 2)                                                # 25 gönderi = 2 paket
            self.assertIn('satis_dili="yok"', MESAJLAR[0][1])
            self.assertIn('karakter_siniri="280"', MESAJLAR[0][1])
            g = {x.id: x for x in s["gonderiler"]}
            k = {i: " | ".join(a for _, a in x.kontrol) for i, x in g.items()}
            self.assertIn("Kanıt gerektiren iddia: 'en iyi'", k["G001"])
            self.assertIn("karakter sınırı aşıldı: 300 / 280", k["G002"])
            self.assertIn("Yasaklı ifade: 'ucuz'", k["G003"])
            self.assertIn("Hashtag sayısı 7", k["G004"])
            self.assertIn("taslak döndürmedi", k["G005"])
            self.assertEqual(k["G007"], "")                                                   # %15 kampanya mesajında var
            self.assertIn("Anma / farkındalık gününde satış ifadesi: indirim", k["G008"])
            self.assertIn("olmayan oran / tutar: %25", k["G011"])
            self.assertEqual(k["G024"], "")                                                   # süren kampanyanın oranı
            self.assertEqual(g["G006"].hashtagler, ["#kahve", "#demlemeipucu"])               # # eklendi, boşluk ve tekrar temizlendi
            self.assertNotIn("G999", g)
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Takvim", "Özet", "Özel Günler", "Uyarılar"])

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(agent.main(["--model-yok", "--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(agent.main(["--model-yok", "--ay", "Kasım"]), 2)
            self.assertEqual(agent.main(["--profil", str(Path(tmp) / "yok.txt"), "--ay", "2026-11"]), 1)
            kotu = Path(tmp) / "p.txt"
            kotu.write_text("Marka: X\nPlatformlar: MySpace=3\n", encoding="utf-8")
            self.assertEqual(agent.main(["--model-yok", "--profil", str(kotu), "--ay", "2026-11"]), 1)


if __name__ == "__main__":
    unittest.main()
