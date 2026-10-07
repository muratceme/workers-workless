"""
Hasar Dosyası Özeti — Workers / Workless AI Agent
Sigorta › Hasar Yönetimi › Hasar Uzmanı

1. Dosya klasöründeki belgeler (.pdf, .docx, .txt) okunur ve kodla türüne göre ayrılır (poliçe, ihbar/beyan,
   kaza tespit tutanağı, ekspertiz raporu, fatura, ehliyet, ruhsat).
2. Kod; plakaları, hasar tarihlerini, poliçe dönemini ve tutarları çıkarır ve belgeler arasında karşılaştırır:
   tarih uyuşmazlığı, sigortalı plakasına benzeyen ama farklı plaka (yazım hatası?), poliçe dönemi dışında veya
   başlangıca yakın hasar, faturanın eksper onaylı tutarı aşması, ihbar gecikmesi ve eksik evrak.
3. Kişi adları (--gizle), telefon, e-posta, TCKN, IBAN maskelenir; plakalar takma adlarla ([PLAKA-A]) tutarlı
   biçimde değiştirilir. Model, kod bulgularını esas alarak karar için tek sayfalık özet yazır.
4. Çıktı: tek sayfalık Markdown özet + Excel (belgeler, kontroller, özet, 'Uzman Onayı').

Kullanım:
    python agent.py                                              # örnek dosyayla dener
    python agent.py --girdi ./dosya_2026_0457 --gizle "Ad Soyad" "Diğer Sürücü"
    python agent.py --girdi ./dosya --evrak evrak_listesi.txt
"""
from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
YAKIN_BASLANGIC_GUN = 30

TURLER = [  # (tür, başlıkta/metinde aranan ifadeler) — sıra önemli
    ("Ekspertiz raporu", ("ekspertiz raporu", "eksper raporu", "ekspertiz")),
    ("Kaza tespit tutanağı", ("kaza tespit tutanağı", "trafik kazası tespit tutanağı", "maddi hasarlı trafik kazası")),
    ("İhbar / beyan", ("hasar ihbar", "ihbar formu", "sigortalı beyanı", "beyan")),
    ("Poliçe", ("poliçe",)),
    ("Fatura", ("fatura",)),
    ("Ehliyet", ("sürücü belgesi", "ehliyet")),
    ("Ruhsat", ("ruhsat", "tescil belgesi")),
    ("Fotoğraf / diğer", ()),
]
VARSAYILAN_EVRAK = ["Poliçe", "İhbar / beyan", "Kaza tespit tutanağı", "Ekspertiz raporu", "Fatura", "Ehliyet", "Ruhsat"]

PLAKA = re.compile(r"\b(0[1-9]|[1-7]\d|8[01])\s?([A-ZÇĞİÖŞÜ]{1,3})\s?(\d{2,5})\b")
TARIH = re.compile(r"\b(\d{1,2})[./](\d{1,2})[./](\d{4})\b")
TUTAR = re.compile(r"(\d{1,3}(?:\.\d{3})*(?:,\d{2})|\d+(?:,\d{2}))\s*(?:TL|₺)")


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower()


def para(s: str) -> Decimal | None:
    try:
        return Decimal(s.replace(".", "").replace(",", "."))
    except InvalidOperation:
        return None


def tarih_al(metin: str, *etiketler: str) -> date | None:
    for et in etiketler:
        m = re.search(re.escape(et) + r"[^\n\d]{0,25}(\d{1,2}[./]\d{1,2}[./]\d{4})", metin, re.I)
        if m:
            g, a, y = TARIH.search(m.group(1)).groups()
            try:
                return date(int(y), int(a), int(g))
            except ValueError:
                return None
    return None


def tutar_al(metin: str, *etiketler: str) -> Decimal | None:
    for et in etiketler:
        m = re.search(r"^[^\n]*" + re.escape(et) + r"[^\n]*?" + TUTAR.pattern, metin, re.I | re.M)
        if m:
            return para(m.group(1))
    return None


def plaka_norm(p: tuple | str) -> str:
    return "".join(p) if isinstance(p, tuple) else re.sub(r"\s", "", p)


# ----------------------------------------------------------------------------
# Belgeler ve kod kontrolleri
# ----------------------------------------------------------------------------

@dataclass
class Belge:
    ad: str
    tur: str
    metin: str
    plakalar: list[str] = field(default_factory=list)
    hasar_tarihi: date | None = None


def tur_bul(ad: str, metin: str) -> str:
    bas = kucuk(ad.replace("_", " ") + "\n" + "\n".join(metin.strip().splitlines()[:3]))
    for tur, ifadeler in TURLER:
        if any(i in bas for i in ifadeler):
            return tur
    govde = kucuk(metin)
    for tur, ifadeler in TURLER:
        if any(i in govde for i in ifadeler):
            return tur
    return "Fotoğraf / diğer"


def belgeleri_oku(klasor: Path) -> tuple[list[Belge], list[str]]:
    belgeler, atlanan = [], []
    for yol in sorted(klasor.iterdir()):
        if yol.suffix.lower() not in belge.DESTEKLENEN:
            atlanan.append(yol.name)
            continue
        try:
            metin = belge.metin_oku(yol)
        except belge.BelgeHatasi as h:
            atlanan.append(f"{yol.name} ({h})")
            continue
        b = Belge(yol.name, tur_bul(yol.stem, metin), metin)
        b.plakalar = list(dict.fromkeys(plaka_norm(p) for p in PLAKA.findall(metin.upper().replace("İ", "I"))))
        b.hasar_tarihi = tarih_al(metin, "olay tarihi", "kaza tarihi", "hasar tarihi", "hasar günü")
        belgeler.append(b)
    return belgeler, atlanan


def benzer(a: str, b: str) -> bool:
    """Aynı uzunlukta tek karakter farkı ya da bir karakter eksik/fazla (yazım hatası)."""
    if a == b:
        return False
    if len(a) == len(b):
        return sum(x != y for x, y in zip(a, b)) == 1
    kisa, uzun = sorted((a, b), key=len)
    return len(uzun) - len(kisa) == 1 and any(uzun[:i] + uzun[i + 1:] == kisa for i in range(len(uzun)))


def kontrol_et(belgeler: list[Belge], evrak_listesi: list[str]) -> dict:
    bulgular: list[tuple[str, str, list[str]]] = []      # (önem, açıklama, kaynak belgeler)
    tur = {b.tur: b for b in belgeler}
    police = tur.get("Poliçe")
    ozet: dict = {}

    # Sigortalı plakası: poliçede ilk plaka, yoksa en çok geçen
    sayac: dict[str, list[str]] = {}
    for b in belgeler:
        for p in b.plakalar:
            sayac.setdefault(p, []).append(b.ad)
    sigortali = (police.plakalar[0] if police and police.plakalar else max(sayac, key=lambda p: len(sayac[p]))) if sayac else None
    ozet["sigortali_plaka"] = sigortali
    if sigortali:
        for p, kaynak in sayac.items():
            if benzer(p, sigortali):
                bulgular.append(("yüksek", f"Plaka uyuşmazlığı: {', '.join(kaynak)} belgesinde '{p}', diğer belgelerde '{sigortali}' "
                                           "(yazım hatası veya farklı araç olabilir)", kaynak + ([police.ad] if police else [])))

    # Hasar tarihi
    tarihler = {b.ad: b.hasar_tarihi for b in belgeler if b.hasar_tarihi and b.tur != "Poliçe"}
    farkli = sorted(set(tarihler.values()))
    ozet["hasar_tarihi"] = max(set(tarihler.values()), key=list(tarihler.values()).count) if tarihler else None
    if len(farkli) > 1:
        bulgular.append(("yüksek", "Hasar tarihi belgelerde farklı: " + "; ".join(f"{ad}: {t:%d.%m.%Y}" for ad, t in tarihler.items()),
                         list(tarihler)))

    # Poliçe dönemi
    if police:
        bas = tarih_al(police.metin, "başlangıç tarihi", "başlangıç")
        bit = tarih_al(police.metin, "bitiş tarihi", "bitiş")
        ozet["police_donemi"] = (bas, bit)
        for ad, t in tarihler.items():
            if bas and bit and not (bas <= t <= bit):
                bulgular.append(("kritik", f"{ad}: hasar tarihi {t:%d.%m.%Y} poliçe dönemi ({bas:%d.%m.%Y}–{bit:%d.%m.%Y}) dışında", [ad, police.ad]))
        t = ozet["hasar_tarihi"]
        if bas and t and 0 <= (t - bas).days <= YAKIN_BASLANGIC_GUN:
            bulgular.append(("orta", f"Hasar, poliçe başlangıcından {(t - bas).days} gün sonra gerçekleşmiş (başlangıca yakın hasar; "
                                     "incelenmesi önerilir)", [police.ad]))

    # İhbar gecikmesi
    ihbar = tur.get("İhbar / beyan")
    if ihbar and ozet["hasar_tarihi"]:
        it = tarih_al(ihbar.metin, "ihbar tarihi", "bildirim tarihi")
        if it:
            ozet["ihbar_gecikmesi"] = (it - ozet["hasar_tarihi"]).days
            if ozet["ihbar_gecikmesi"] > 5:
                bulgular.append(("orta", f"İhbar, olaydan {ozet['ihbar_gecikmesi']} gün sonra yapılmış", [ihbar.ad]))

    # Tutarlar: eksper (KDV hariç) ve fatura
    eksper = tur.get("Ekspertiz raporu")
    fatura = tur.get("Fatura")
    e_tutar = tutar_al(eksper.metin, "toplam") if eksper else None
    f_ara = tutar_al(fatura.metin, "ara toplam", "kdv hariç", "matrah") if fatura else None
    f_genel = tutar_al(fatura.metin, "genel toplam", "ödenecek tutar", "toplam tutar") if fatura else None
    ozet.update({"eksper_tutar": e_tutar, "fatura_ara": f_ara, "fatura_genel": f_genel})
    if e_tutar is not None and f_ara is not None and f_ara != e_tutar:
        fark = f_ara - e_tutar
        bulgular.append(("yüksek" if fark > 0 else "düşük",
                         f"Fatura KDV hariç tutarı ({f_ara:,.2f} TL) eksper onaylı tutardan ({e_tutar:,.2f} TL) "
                         f"{abs(fark):,.2f} TL {'fazla' if fark > 0 else 'az'}".replace(",", "X").replace(".", ",").replace("X", "."),
                         [fatura.ad, eksper.ad]))

    # Eksik evrak
    mevcut = {b.tur for b in belgeler}
    ozet["eksik_evrak"] = [e for e in evrak_listesi if e not in mevcut]
    for e in ozet["eksik_evrak"]:
        bulgular.append(("orta", f"Eksik evrak: {e}", []))
    return {"bulgular": bulgular, "ozet": ozet}


# ----------------------------------------------------------------------------
# Maskeleme
# ----------------------------------------------------------------------------

def takma_adlar(belgeler: list[Belge], sigortali: str | None, terimler: list[str]) -> dict[str, str]:
    harita = {}
    plakalar = list(dict.fromkeys(([sigortali] if sigortali else []) + [p for b in belgeler for p in b.plakalar]))
    for i, p in enumerate(plakalar):
        harita[p] = f"[PLAKA-{chr(65 + i) if i < 26 else i}]"
    for i, t in enumerate(sorted({t.strip() for t in terimler if t.strip()}, key=len, reverse=True), 1):
        harita[t] = f"[GİZLİ-{i}]"
    return harita


def maskele(metin: str, harita: dict[str, str]) -> str:
    metin = llm.maskele(metin)
    def plaka_degis(m):
        return harita.get(plaka_norm(m.groups()), m.group(0))
    metin = PLAKA.sub(plaka_degis, metin)
    for t, takma in harita.items():
        if not PLAKA.fullmatch(t):                       # plakalar yukarıda değiştirildi
            metin = re.sub(re.escape(t), takma, metin, flags=re.I)
    return metin


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

SEMA = {
    "type": "object",
    "properties": {
        "olay_ozeti": {"type": "string"},
        "kusur_ve_rucu": {"type": "string"},
        "teminat_notu": {"type": "string"},
        "tutarsizliklar": {"type": "array", "items": {
            "type": "object",
            "properties": {"aciklama": {"type": "string"}, "kaynaklar": {"type": "array", "items": {"type": "string"}}},
            "required": ["aciklama", "kaynaklar"], "additionalProperties": False}},
        "acik_sorular": {"type": "array", "items": {"type": "string"}},
        "onerilen_adimlar": {"type": "array", "items": {"type": "string"}},
        "karar_ozeti": {"type": "string"},
    },
    "required": ["olay_ozeti", "kusur_ve_rucu", "teminat_notu", "tutarsizliklar", "acik_sorular", "onerilen_adimlar", "karar_ozeti"],
    "additionalProperties": False,
}


def ozetle(belgeler: list[Belge], kontrol: dict, harita: dict[str, str]) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    kk = [f"- [{onem}] {maskele(acik, harita)} (kaynak: {', '.join(k) or '-'})" for onem, acik, k in kontrol["bulgular"]]
    o = kontrol["ozet"]
    for ad, deger in (("Eksper onaylı tutar (KDV hariç)", o.get("eksper_tutar")), ("Fatura ara toplam (KDV hariç)", o.get("fatura_ara")),
                      ("Fatura genel toplam (KDV dahil)", o.get("fatura_genel"))):
        if deger is not None:
            kk.append(f"- {ad}: {deger} TL")
    mesaj = "\n".join([
        "<belgeler>",
        *[f'<belge ad="{b.ad}" tur="{b.tur}">\n{maskele(b.metin, harita)}\n</belge>' for b in belgeler],
        "</belgeler>",
        "<kod_kontrolleri>", *(kk or ["- Tutarsızlık bulunmadı"]), "</kod_kontrolleri>",
    ])
    yanit = llm.json_iste(sistem, mesaj, SEMA)
    adlar = {b.ad for b in belgeler}
    for t in yanit["tutarsizliklar"]:
        gecersiz = [k for k in t["kaynaklar"] if k not in adlar]
        t["kaynaklar"] = [k for k in t["kaynaklar"] if k in adlar]
        if gecersiz:
            t["aciklama"] += " [kaynak doğrulanamadı]"
    return yanit


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
ONEM_DOLGU = {"kritik": PatternFill("solid", fgColor="F8C4C1"), "yüksek": PatternFill("solid", fgColor="FDE2E1"),
              "orta": PatternFill("solid", fgColor="FFF4CE"), "düşük": PatternFill("solid", fgColor="E3F5E1")}
ONAY_DOLGU = PatternFill("solid", fgColor="FFF4CE")
UST = Alignment(vertical="top", wrap_text=True)


def tl(x) -> str:
    return "—" if x is None else f"{x:,.2f} TL".replace(",", "X").replace(".", ",").replace("X", ".")


def geri_ac(metin: str, harita: dict[str, str]) -> str:
    """Rapor şirket içinde kalır: takma adları gerçek değerlere geri çevir (plakalar boşluklu yazılır)."""
    for gercek, takma in harita.items():
        m = PLAKA.fullmatch(gercek)
        metin = metin.replace(takma, " ".join(m.groups()) if m else gercek)
    return metin


def rapor_yaz(cikti: Path, klasor: Path, belgeler: list[Belge], atlanan: list[str], kontrol: dict, ozet: dict, harita: dict) -> Path:
    o = kontrol["ozet"]
    g = lambda s: geri_ac(s, harita)  # noqa: E731
    md = [f"# Hasar dosyası özeti — {klasor.name}", "",
          f"**Hasar tarihi:** {o['hasar_tarihi']:%d.%m.%Y}" if o.get("hasar_tarihi") else "**Hasar tarihi:** —",
          f" · **Sigortalı plaka:** {g(harita.get(o['sigortali_plaka'], '—')) if o.get('sigortali_plaka') else '—'}"
          f" · **Eksper (KDV hariç):** {tl(o.get('eksper_tutar'))} · **Fatura (KDV dahil):** {tl(o.get('fatura_genel'))}", "",
          "## Olay", g(ozet["olay_ozeti"]), "", "## Kusur ve rücu", g(ozet["kusur_ve_rucu"]), "",
          "## Teminat değerlendirmesi (taslak)", g(ozet["teminat_notu"]), "", "## Tutarsızlıklar ve kontrol bulguları"]
    for onem, acik, kaynak in kontrol["bulgular"]:
        md.append(f"- **[{onem}]** {acik}" + (f" _({', '.join(kaynak)})_" if kaynak else "") + " — _kod kontrolü_")
    for t in ozet["tutarsizliklar"]:
        md.append(f"- {g(t['aciklama'])}" + (f" _({', '.join(t['kaynaklar'])})_" if t["kaynaklar"] else ""))
    md += ["", "## Açık sorular", *[f"- {g(s)}" for s in ozet["acik_sorular"]], "", "## Önerilen adımlar",
           *[f"{i}. {g(s)}" for i, s in enumerate(ozet["onerilen_adimlar"], 1)], "", "## Karar özeti (öneri taslağı)",
           g(ozet["karar_ozeti"]), "", "---",
           f"_Belgeler: {', '.join(f'{b.ad} ({b.tur})' for b in belgeler)}_" + (f"  \n_Okunamayan: {', '.join(atlanan)}_" if atlanan else ""),
           "_Yapay zekâ destekli özet; kod kontrolleri kesin, değerlendirmeler taslaktır. Karar hasar uzmanına aittir._"]
    cikti.parent.mkdir(parents=True, exist_ok=True)
    md_yol = cikti.with_suffix(".md")
    md_yol.write_text("\n".join(md), encoding="utf-8")

    wb = Workbook()
    ws = wb.active
    ws.title = "Kontroller"
    ws.append(["Önem", "Bulgu", "Kaynak Belgeler", "Kaynak", "Uzman Onayı"])
    for h in ws[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for onem, acik, kaynak in kontrol["bulgular"]:
        ws.append([onem, acik, ", ".join(kaynak), "Kod", ""])
        ws.cell(ws.max_row, 1).fill = ONEM_DOLGU[onem]
        ws.cell(ws.max_row, 5).fill = ONAY_DOLGU
    for t in ozet["tutarsizliklar"]:
        ws.append(["model", g(t["aciklama"]), ", ".join(t["kaynaklar"]), "Model", ""])
        ws.cell(ws.max_row, 5).fill = ONAY_DOLGU
    for j, w in enumerate((10, 100, 40, 8, 14), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for satir in ws.iter_rows():
        for h in satir:
            h.alignment = UST
    b = wb.create_sheet("Belgeler")
    b.append(["Belge", "Tür (kod)", "Plakalar", "Hasar Tarihi", "Karakter"])
    for h in b[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for x in belgeler:
        b.append([x.ad, x.tur, ", ".join(x.plakalar), x.hasar_tarihi, len(x.metin)])
        b.cell(b.max_row, 4).number_format = "DD.MM.YYYY"
    for ad in atlanan:
        b.append([ad, "okunamadı"])
    for j, w in enumerate((32, 22, 30, 13, 10), 1):
        b.column_dimensions[get_column_letter(j)].width = w
    s = wb.create_sheet("Özet", 0)
    for satir in [["Hasar dosyası özeti", klasor.name], [], ["Olay", g(ozet["olay_ozeti"])], ["Kusur ve rücu", g(ozet["kusur_ve_rucu"])],
                  ["Teminat (taslak)", g(ozet["teminat_notu"])], ["Karar özeti (öneri)", g(ozet["karar_ozeti"])],
                  ["Açık sorular", "\n".join(f"• {g(x)}" for x in ozet["acik_sorular"])],
                  ["Önerilen adımlar", "\n".join(f"{i}. {g(x)}" for i, x in enumerate(ozet["onerilen_adimlar"], 1))],
                  ["Model", llm.kullanim_ozeti()], ["Oluşturulma", datetime.now().strftime("%d.%m.%Y %H:%M")]]:
        s.append(satir)
    s["A1"].font = Font(bold=True, size=13)
    s.column_dimensions["A"].width = 22
    s.column_dimensions["B"].width = 120
    for satir in s.iter_rows():
        for h in satir:
            h.alignment = UST
    wb.save(cikti)
    return md_yol


def calistir(klasor: Path, cikti: Path, terimler: list[str] | None = None, evrak_yolu: Path | None = None, evet: bool = False) -> dict:
    if not klasor.is_dir():
        raise llm.LLMHatasi(f"{klasor}: klasör bulunamadı (hasar dosyasının belgelerini bir klasörde toplayın).")
    belgeler, atlanan = belgeleri_oku(klasor)
    if not belgeler:
        raise llm.LLMHatasi(f"{klasor}: okunabilir belge yok.")
    evrak = VARSAYILAN_EVRAK if not evrak_yolu else \
        [x.strip() for x in evrak_yolu.read_text(encoding="utf-8").splitlines() if x.strip() and not x.startswith("#")]
    kontrol = kontrol_et(belgeler, evrak)
    print(f"[OK] {len(belgeler)} belge · kod kontrolü: {len(kontrol['bulgular'])} bulgu")
    for onem, acik, _ in kontrol["bulgular"]:
        print(f"[!] [{onem}] {acik}")
    harita = takma_adlar(belgeler, kontrol["ozet"].get("sigortali_plaka"), terimler or [])
    llm.onay_al(f"{len(belgeler)} belgenin metni (plakalar takma adlı; telefon/e-posta/TCKN/IBAN ve {len(terimler or [])} "
                "gizli terim maskeli) özet için gönderilecek.", evet)
    ozet = ozetle(belgeler, kontrol, harita)
    md = rapor_yaz(cikti, klasor, belgeler, atlanan, kontrol, ozet, harita)
    return {"belgeler": belgeler, "kontrol": kontrol, "ozet": ozet, "md": md, "harita": harita}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Hasar dosyasındaki belgelerden karar için tek sayfalık özet hazırlar.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "dosya_2026_0457", help="Hasar dosyası klasörü (.pdf/.docx/.txt)")
    p.add_argument("--gizle", nargs="*", default=[], metavar="TERİM", help="Modele gönderilmeden maskelenecek kişi adları")
    p.add_argument("--evrak", type=Path, help="Gerekli evrak listesi (.txt, satır başına bir belge türü)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "hasar_ozeti.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    if a.girdi == BURASI / "ornek_veri" / "dosya_2026_0457":
        a.gizle = a.gizle or ["Mehmet Örnek", "Ali Deneme"]
    try:
        s = calistir(a.girdi, a.cikti, a.gizle, a.evrak, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] Özet: {s['md'].resolve()}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
