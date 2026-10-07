"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

DOSYA = KLASOR / "ornek_veri" / "dosya_2026_0457"
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return {"olay_ozeti": "[PLAKA-A] plakalı araca [PLAKA-B] arkadan çarptı.", "kusur_ve_rucu": "Kusur [PLAKA-B] sürücüsü [GİZLİ-2].",
            "teminat_notu": "Taslak.", "acik_sorular": ["Fatura plakası neden farklı?"], "onerilen_adimlar": ["Ehliyet ve ruhsat isteyin."],
            "karar_ozeti": "Eksikler giderilince değerlendirilebilir.",
            "tutarsizliklar": [{"aciklama": "Ek bulgu", "kaynaklar": ["servis_faturasi.txt", "uydurma.pdf"]}]}


class KontrolTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.belgeler, _ = agent.belgeleri_oku(DOSYA)
        cls.k = agent.kontrol_et(cls.belgeler, agent.VARSAYILAN_EVRAK)

    def test_tur_siniflandirma(self):
        self.assertEqual({b.ad: b.tur for b in self.belgeler}, {
            "ekspertiz_raporu.txt": "Ekspertiz raporu", "hasar_ihbar_beyani.txt": "İhbar / beyan",
            "kaza_tespit_tutanagi.txt": "Kaza tespit tutanağı", "police_ozeti.txt": "Poliçe", "servis_faturasi.txt": "Fatura"})

    def test_bulgular(self):
        metin = "\n".join(b[1] for b in self.k["bulgular"])
        self.assertIn("'34ABC128'", metin)                               # fatura plakası farklı
        self.assertNotIn("06XYZ789", metin)                              # karşı araç uyuşmazlık sayılmaz
        self.assertIn("15.09.2026", metin)                               # eksper hasar tarihi farklı
        self.assertIn("13 gün sonra", metin)                             # poliçe başlangıcına yakın
        self.assertIn("2.250,00 TL fazla", metin)                        # 51.000 − 48.750
        self.assertEqual(self.k["ozet"]["eksik_evrak"], ["Ehliyet", "Ruhsat"])
        self.assertEqual(self.k["ozet"]["hasar_tarihi"], date(2026, 9, 14))   # çoğunluk
        self.assertEqual((self.k["ozet"]["eksper_tutar"], self.k["ozet"]["fatura_genel"]), (D("48750.00"), D("61200.00")))

    def test_police_donemi_disi(self):
        b = [agent.Belge("p.txt", "Poliçe", "Poliçe Başlangıç Tarihi: 01.09.2026\nBitiş Tarihi: 01.09.2027"),
             agent.Belge("t.txt", "Kaza tespit tutanağı", "", hasar_tarihi=date(2026, 8, 30))]
        k = agent.kontrol_et(b, [])
        self.assertTrue(any(o == "kritik" and "dışında" in a for o, a, _ in k["bulgular"]))

    def test_benzer(self):
        self.assertTrue(agent.benzer("34ABC123", "34ABC128"))
        self.assertTrue(agent.benzer("34ABC123", "34AB123"))
        self.assertFalse(agent.benzer("34ABC123", "06XYZ789"))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_uctan_uca(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "h.xlsx"
            s = agent.calistir(DOSYA, cikti, ["Mehmet Örnek", "Ali Deneme"], evet=True)
            m = MESAJLAR[0]
            for gizli in ("Mehmet Örnek", "Ali Deneme", "34 ABC 123", "06 XYZ 789", "0533 222 33 44"):
                self.assertNotIn(gizli, m)
            self.assertIn("[PLAKA-A]", m)
            self.assertIn("<kod_kontrolleri>", m)
            md = s["md"].read_text(encoding="utf-8")
            self.assertIn("34 ABC 123", md)                               # rapor şirket içi: takma adlar geri açılır
            self.assertNotIn("[PLAKA-A]", md)
            self.assertIn("Ali Deneme", md)
            self.assertIn("[kaynak doğrulanamadı]", md)                    # uydurma belge adı işaretlenir
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Özet", "Kontroller", "Belgeler"])
            self.assertEqual(wb["Kontroller"].cell(1, 5).value, "Uzman Onayı")

    def test_klasor_yok(self):
        with self.assertRaises(llm.LLMHatasi):
            agent.calistir(Path("olmayan_klasor"), Path("x.xlsx"), evet=True)


if __name__ == "__main__":
    unittest.main()
