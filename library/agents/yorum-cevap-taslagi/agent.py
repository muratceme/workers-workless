"""
Yorum Cevap Taslağı — Workers / Workless AI Agent
Misafir İlişkileri › Misafir İlişkileri Müdürü

1. Kod otelin üslup rehberini okur (imza satırı "İmza: …"), yorumları ve misafir adlarını hazırlar.
2. Model her yorum için yorumdaki noktaları (olumlu / olumsuz, alıntıyla) çıkarır ve yorumun dilinde, rehberdeki üslupla
   kişiselleştirilmiş, herkese açık bir cevap yazar; hangi noktalara değindiğini belirtir.
3. Kod denetler:
   - Alıntılar yorumda var mı; her olumsuz noktaya değinilmiş mi
   - Tazminat / indirim / iade / ücretsiz konaklama teklifi var mı (--teklif-izni yoksa yasak)
   - Misafiri suçlayan ifade var mı
   - Oda / rezervasyon numarası olabilecek 3+ haneli sayı ya da yorumda olmayan sayı var mı
   - İmza ve {AD} hitabı var mı; --azami karakteri aşıyor mu
   - İki cevap birbirine çok benziyor mu (kalıp cevap)
   Sorun varsa yalnız o yorumun cevabı modele geri gönderilir (en fazla --tur kez).
4. Cevaplar yayımlanmaz; Excel'de "Onay" sütunuyla listelenir.

Kullanım:
    python agent.py                                               # örnek yorumlar + üslup rehberi
    python agent.py --girdi yorumlar.xlsx --rehber uslup.txt --azami 1000
"""
from __future__ import annotations

import argparse
import csv
import difflib
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
PAKET = 15
SEMA = {
    "type": "object",
    "properties": {"cevaplar": {"type": "array", "items": {"type": "object", "properties": {
        "id": S, "dil": S,
        "noktalar": {"type": "array", "items": {"type": "object", "properties": {"nokta": S, "duygu": {"type": "string", "enum": ["Olumlu", "Olumsuz"]}, "alinti": S},
                                                "required": ["nokta", "duygu", "alinti"], "additionalProperties": False}},
        "cevap": S, "degindigi_noktalar": {"type": "array", "items": {"type": "integer"}}, "ic_not": S},
        "required": ["id", "dil", "noktalar", "cevap", "degindigi_noktalar", "ic_not"], "additionalProperties": False}}},
    "required": ["cevaplar"],
    "additionalProperties": False,
}
TEKLIF = re.compile(r"(indirim|iade|ücretsiz|ucretsiz|tazminat|hediye konaklama|bedava|refund|discount|compensat|free (night|stay)|voucher|"
                    r"rabatt|erstatt|kostenlos|gutschein|скидк|возврат|бесплатн)", re.I)
SUCLAMA = re.compile(r"(sizin hatanız|kendi hatanız|kurallara uymadınız|haksızsınız|you should have|your fault|you failed to|ihre schuld|sie hätten)", re.I)
BENZERLIK = 0.85


def rehber_oku(yol: Path) -> tuple[str, str]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: üslup rehberi bulunamadı.")
    metin = belge.metin_oku(yol)
    m = re.search(r"^\s*İmza\s*:\s*(.+)$", metin, re.M | re.I)
    return metin, (m.group(1).strip() if m else "")


def yorumlari_oku(yol: Path) -> list[dict]:
    r = kayitlar(yol, {"id": ("Yorum No", "No"), "tarih": ("Tarih",), "platform": ("Platform", "Kanal", "Site"), "puan": ("Puan", "Rating"),
                       "ad": ("Misafir", "Misafir Adı", "Kullanıcı", "Ad"), "yorum": ("Yorum", "Metin", "Değerlendirme")}, ("yorum",))
    return [{"id": str(x["id"] or f"YR{i}"), "tarih": tarih_coz(x["tarih"]), "platform": str(x["platform"]), "puan": str(x["puan"]), "ad": str(x["ad"]).split()[0] if str(x["ad"]).strip() else "",
             "yorum": str(x["yorum"]).strip()} for i, x in enumerate(r, 1) if str(x["yorum"]).strip()]


def denetle(yorum: list[dict], sonuc: dict, rehber: str, imza: str, azami: int, teklif_izni: bool) -> list[str]:
    sorun, gorulen = [], []
    for y in yorum:
        c = sonuc.get(y["id"])
        if not c:
            sorun.append(f"{y['id']}: cevap yazılmadı")
            continue
        ad, cv = y["id"], c["cevap"]
        for i, n in enumerate(c["noktalar"], 1):
            if not metinde_var(n["alinti"], y["yorum"]):
                sorun.append(f"{ad}: nokta {i} alıntısı yorumda yok: '{n['alinti'][:50]}'")
        eksik = [i for i, n in enumerate(c["noktalar"], 1) if n["duygu"] == "Olumsuz" and i not in c["degindigi_noktalar"]]
        if eksik:
            sorun.append(f"{ad}: olumsuz nokta {eksik} cevapta ele alınmamış")
        if not c["noktalar"]:
            sorun.append(f"{ad}: yorumdaki noktalar çıkarılmamış")
        if not teklif_izni and TEKLIF.search(cv):
            sorun.append(f"{ad}: cevapta tazminat / indirim / iade / ücretsiz teklif ifadesi var ('{TEKLIF.search(cv).group(0)}')")
        if SUCLAMA.search(cv):
            sorun.append(f"{ad}: misafiri suçlayan ifade ('{SUCLAMA.search(cv).group(0)}')")
        uzun = [x for x in sayilar(cv) if re.fullmatch(r"\d{3,}", x)]
        if uzun:
            sorun.append(f"{ad}: cevapta oda / rezervasyon numarası olabilecek sayı: {', '.join(sorted(uzun))}")
        fazla = sorted(sayilar(cv) - sayilar(y["yorum"] + " " + rehber) - set(uzun))
        if fazla:
            sorun.append(f"{ad}: yorumda ve rehberde olmayan sayı: {', '.join(fazla)}")
        if imza and katla(imza) not in katla(cv):
            sorun.append(f"{ad}: imza ('{imza}') yok")
        if "{AD}" not in cv:
            sorun.append(f"{ad}: hitapta {{AD}} yer tutucusu yok")
        if len(cv) > azami:
            sorun.append(f"{ad}: cevap {len(cv)} karakter, sınır {azami}")
        for oid, ocv in gorulen:
            if difflib.SequenceMatcher(None, katla(ocv), katla(cv)).ratio() > BENZERLIK:
                sorun.append(f"{ad}: cevap {oid} ile neredeyse aynı; yorumdaki ayrıntılara göre kişiselleştirin")
                break
        gorulen.append((ad, cv))
    return sorun


def uret(yorum: list[dict], rehber: str, azami: int, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    sonuc, hedef = {}, yorum
    if duzeltme is not None:  # düzeltmede yalnız sorunlu cevaplar yeniden yazılır
        sorunlu = {s.split(":")[0] for s in duzeltme}
        sonuc.update({a: b for a, b in (onceki or {}).items() if a not in sorunlu})
        hedef = [y for y in yorum if y["id"] in sorunlu]
    for j in range(0, len(hedef), PAKET):
        parca = hedef[j:j + PAKET]
        ids = {y["id"] for y in parca}
        satir = "\n".join(f'<yorum id="{y["id"]}" platform="{y["platform"]}" puan="{y["puan"]}">{llm.maskele(y["yorum"])}</yorum>' for y in parca)
        d = [s for s in (duzeltme or []) if s.split(":")[0] in ids]
        mesaj = (f"<uslup_rehberi>\n{rehber}\n</uslup_rehberi>\n<azami_karakter>{azami}</azami_karakter>\n<yorumlar>\n{satir}\n</yorumlar>"
                 + duzeltme_blogu(d, [v for a, v in (onceki or {}).items() if a in ids]))
        for c in llm.json_iste(sistem, mesaj, SEMA).get("cevaplar", []):
            if c.get("id") in ids:
                sonuc[c["id"]] = c
    return sonuc


def doldur(cevap: str, ad: str, dil: str) -> str:
    return cevap.replace("{AD}", ad or {"tr": "Değerli misafirimiz", "de": "Gast"}.get(dil.lower()[:2], "Guest"))


def rapor_yaz(cikti: Path, yorum, sonuc, kalan, tur) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Cevaplar"
    baslik(ws, ["Yorum No", "Tarih", "Platform", "Puan", "Misafir", "Yorum", "Dil", "Olumlu", "Olumsuz", "Cevap Taslağı", "Karakter", "İç Not", "Onay"],
           (8, 11, 11, 6, 10, 50, 5, 34, 34, 70, 9, 30, 8))
    for y in yorum:
        c = sonuc.get(y["id"])
        if not c:
            ws.append([y["id"], y["tarih"], y["platform"], y["puan"], y["ad"], y["yorum"], "", "", "", "AI yanıt vermedi"])
            continue
        cv = doldur(c["cevap"], y["ad"], c["dil"])
        ws.append([y["id"], y["tarih"], y["platform"], y["puan"], y["ad"], y["yorum"], c["dil"],
                   "\n".join(n["nokta"] for n in c["noktalar"] if n["duygu"] == "Olumlu"), "\n".join(n["nokta"] for n in c["noktalar"] if n["duygu"] == "Olumsuz"),
                   cv, len(cv), c["ic_not"], ""])
        ws.cell(ws.max_row, 2).number_format = "DD.MM.YYYY"
        boya(ws, [8, 9, 10], AI_DOLGU)
        if any(n["duygu"] == "Olumsuz" for n in c["noktalar"]):
            boya(ws, [9], KIRMIZI)
        boya(ws, [13], SARI)
    ws.freeze_panes = "B2"
    son_not(ws, "Önemli: Cevaplar yayımlanmadı. Okuyup onayladıktan sonra platformda yayımlayın; olumsuz noktaları ilgili departmana iletin.", kalan, tur)
    hizala(ws)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(girdi: Path, rehber_yolu: Path, cikti: Path, azami: int = 1200, teklif_izni: bool = False, tur: int = 2, evet: bool = False) -> dict:
    yorum = yorumlari_oku(girdi)
    rehber, imza = rehber_oku(rehber_yolu)
    if not yorum:
        raise llm.LLMHatasi(f"{girdi.name}: yorum yok.")
    print(f"[OK] {len(yorum)} yorum · imza: {imza or '(rehberde yok)'}")
    llm.onay_al(f"{len(yorum)} yorum metni (kullanıcı adı hariç; telefon / e-posta maskeli) ve üslup rehberi gönderilecek.", evet)
    sonuc, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(yorum, rehber, azami, d, o), lambda x: denetle(yorum, x, rehber, imza, azami, teklif_izni), tur)
    rapor_yaz(cikti, yorum, sonuc, kalan, yapilan)
    return {"yorumlar": yorum, "sonuc": sonuc, "kalan": kalan, "tur": yapilan, "imza": imza}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    p = argparse.ArgumentParser(description="Misafir yorumlarına otelin üslubuna uygun, kişiselleştirilmiş cevap taslakları hazırlar (yayımlamaz).")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "yorumlar.csv",
                   help="Yorumlar (.xlsx/.csv): Yorum; Yorum No, Tarih, Platform, Puan, Misafir isteğe bağlı")
    p.add_argument("--rehber", type=Path, default=BURASI / "ornek_veri" / "uslup_rehberi.txt", help="Üslup rehberi (.txt/.md/.docx); 'İmza: …' satırı")
    p.add_argument("--azami", type=int, default=1200, help="Cevap için en fazla karakter (varsayılan 1200; platform sınırına göre ayarlayın)")
    p.add_argument("--teklif-izni", action="store_true", help="Tazminat / indirim / iade ifadelerine izin ver (varsayılan: yasak)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "yorum_cevaplari.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.rehber, a.cikti, a.azami, a.teklif_izni, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(s['sonuc'])}/{len(s['yorumlar'])} cevap taslağı (yayımlanmadı)")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
