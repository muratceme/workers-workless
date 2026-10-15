"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

YORUM = KLASOR / "ornek_veri" / "yorumlar.csv"
REHBER = KLASOR / "ornek_veri" / "uslup_rehberi.txt"
IMZA = "Misafir İlişkileri Ekibi, Örnek Otel"
MESAJLAR = []
P, O = "Olumlu", "Olumsuz"
DOGRU = {
    "YR-1": ("tr", [("Kahvaltı", P, "Kahvaltı çok zengindi"), ("Resepsiyon", P, "Mert Bey çok yardımcı oldu")], [1, 2],
             "Merhaba {AD}, yerel peynirlerimizi beğenmenize çok sevindik! Güzel sözlerinizi Mert Bey'e de ilettik. Sizi yeniden ağırlamayı dileriz.\n" + IMZA),
    "YR-2": ("en", [("Location", P, "Great location near the old town"), ("Street noise", O, "the room was noisy because of the street"),
                    ("Air conditioning", O, "the air conditioning was not cooling well")], [1, 2, 3],
             "Dear {AD}, thank you for highlighting our location close to the old town. We are sorry about the street noise and the air conditioning; "
             "both points have been shared with our housekeeping and technical teams. We hope to welcome you again.\n" + IMZA),
    "YR-3": ("tr", [("Temizlik", O, "banyoda saç vardı"), ("İlgisizlik", O, "kimse ilgilenmedi"), ("Fiyat-değer", O, "Bu fiyata kabul edilemez")], [1, 2, 3],
             "Sayın {AD}, banyodaki temizlik sorunu ve şikâyetinize zamanında ilgi gösterilmemesi için içtenlikle özür dileriz. Bu deneyim beklediğimiz "
             "standartta değil; yorumunuzu kat hizmetleri ve ön büro yöneticilerimizle paylaştık. Konuyu sizinle doğrudan görüşmek isteriz: "
             "misafir@ornekotel.example\n" + IMZA),
    "YR-4": ("de", [("Personal", P, "freundliches Personal"), ("Frühstück", O, "um 10:30 war schon alles weg")], [1, 2],
             "Liebe {AD}, vielen Dank für Ihre freundlichen Worte über unser Team. Es tut uns leid, dass um 10:30 beim Frühstück nicht mehr alles verfügbar war; "
             "wir haben Ihr Feedback an die Küche weitergegeben.\n" + IMZA),
    "YR-5": ("tr", [("Konum", P, "Konum mükemmel"), ("Oda", P, "odalar ferah")], [1, 2],
             "Merhaba {AD}, ferah odalarımızı ve konumumuzu beğendiğiniz için teşekkür ederiz. Tekrar gelmenizi sabırsızlıkla bekliyoruz!\n" + IMZA),
}


def cevaplar(dogru=True) -> dict:
    d = {i: {"id": i, "dil": dl, "noktalar": [{"nokta": a, "duygu": b, "alinti": c} for a, b, c in n], "cevap": cv, "degindigi_noktalar": dg, "ic_not": ""}
         for i, (dl, n, dg, cv) in DOGRU.items()}
    if not dogru:
        d["YR-3"]["cevap"] = d["YR-3"]["cevap"].replace("banyodaki", "412 numaralı odanızın banyosundaki").replace("Konuyu", "Bir sonraki konaklamanızda indirim sunmak ve konuyu")
        d["YR-2"]["degindigi_noktalar"] = [1, 2]
    return d


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    d = cevaplar("<duzeltme>" in kullanici)
    return {"cevaplar": [d[i] for i in re.findall(r'<yorum id="([^"]+)"', kullanici)]}


class KodTesti(unittest.TestCase):
    def test_rehber(self):
        r, imza = agent.rehber_oku(REHBER)
        self.assertEqual(imza, IMZA)
        self.assertEqual([y["ad"] for y in agent.yorumlari_oku(YORUM)], ["Selin", "Mark", "Hakan", "Julia", "Emre"])

    def test_denetim(self):
        y, (r, imza) = agent.yorumlari_oku(YORUM), agent.rehber_oku(REHBER)
        self.assertEqual(agent.denetle(y, cevaplar(), r, imza, 1200, False), [])
        s = agent.denetle(y, cevaplar(False), r, imza, 1200, False)
        self.assertEqual(len(s), 3)
        self.assertTrue(any("412" in x for x in s) and any("indirim" in x for x in s) and any("YR-2: olumsuz nokta [3]" in x for x in s))
        self.assertEqual(len(agent.denetle(y, cevaplar(False), r, imza, 1200, True)), 2)
        d = cevaplar()
        d["YR-5"]["cevap"] = d["YR-1"]["cevap"]
        d["YR-4"]["cevap"] = d["YR-4"]["cevap"].replace(IMZA, "") + " Sie hätten früher kommen sollen."
        s = agent.denetle(y, d, r, imza, 200, False)
        self.assertTrue(any("YR-5: cevap YR-1 ile neredeyse aynı" in x for x in s) and any("YR-4: imza" in x for x in s)
                        and any("suçlayan" in x for x in s) and any("sınır 200" in x for x in s))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "y.xlsx"
            s = agent.calistir(YORUM, REHBER, cikti, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            self.assertNotIn("Selin", MESAJLAR[0])
            self.assertEqual(re.findall(r'<yorum id="([^"]+)"', MESAJLAR[1]), ["YR-2", "YR-3"])
            ws = load_workbook(cikti)["Cevaplar"]
            self.assertTrue(ws.cell(2, 10).value.startswith("Merhaba Selin,"))
            self.assertEqual(ws.cell(4, 9).value.count("\n"), 2)

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(YORUM, REHBER, Path(tmp) / "y.xlsx", tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 3)
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--rehber", str(REHBER.with_name("yok.txt")), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
