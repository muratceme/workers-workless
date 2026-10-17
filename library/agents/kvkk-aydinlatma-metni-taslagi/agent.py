"""
KVKK Aydınlatma Metni Taslağı — Gerisi Bende AI Agent
Hukuk › Avukat

1. Kod veri işleme envanterini (kategori, veriler, amaç, hukuki sebep, toplama yöntemi, alıcı grubu, aktarım amacı, yurt dışı
   aktarım) ve veri sorumlusu bilgilerini okur; satırları E1, E2 … diye numaralar.
2. Kod envanteri 6698 sayılı KVKK'ya göre ön kontrolden geçirir:
   - Hukuki sebep md. 5 (5/1 açık rıza, 5/2-a … f) ya da md. 6/3 (a … g) bentlerinden biri mi
   - Özel nitelikli veri (sağlık, biyometrik, din, ceza mahkûmiyeti …) md. 5 sebebine dayandırılmış mı
   - Yurt dışı aktarımda md. 9 mekanizması (yeterlilik kararı, standart sözleşme, bağlayıcı şirket kuralları, taahhütname,
     arızi aktarım) yazılmış mı
3. Model aydınlatma metninin amaç, hukuki sebep, aktarım ve toplama yöntemi bölümlerini yazar.
4. Kod denetler:
   - Her envanter satırı bir bölümde ele alınmış mı; bölüm, dayandığı satırların hukuki sebebini açıkça yazıyor mu
   - Envanterde olmayan bir hukuki sebep (ör. açık rıza) metinde geçiyor mu
   - Her alıcı grubu ve yurt dışı aktarım (mekanizmasıyla) yazılmış mı; her toplama yöntemi geçiyor mu
   - "vb.", "ve benzeri", "sınırlı olmamak üzere" gibi genel ifadeler; aydınlatma metnine açık rıza beyanı eklenmiş mi
   Sorun varsa modele geri gönderilir (en fazla --tur kez).
5. Kod veri sorumlusu kimliğini, md. 11 haklarını ve başvuru yöntemlerini aynen ekler.

Kullanım:
    python agent.py                                              # örnek (kurgusal) e-ticaret sitesi
    python agent.py --envanter envanter.xlsx --sorumlu sirket.txt
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
KOLONLAR = {
    "kategori": ("Veri Kategorisi", "Kategori"), "veriler": ("Veriler", "Kişisel Veriler"), "amac": ("İşleme Amacı", "Amaç"),
    "sebep": ("Hukuki Sebep", "Hukuki Dayanak"), "yontem": ("Toplama Yöntemi", "Yöntem"), "alici": ("Alıcı Grubu", "Aktarılan", "Alıcılar"),
    "aktarim_amaci": ("Aktarım Amacı",), "yurtdisi": ("Yurt Dışı Aktarım", "Yurt Dışı"),
}
SORUMLU = {"unvan": "Unvan", "mersis": "MERSİS", "adres": "Adres", "kep": "KEP", "eposta": "E-posta", "grup": "İlgili Kişi Grubu"}
# KVKK md. 5 ve md. 6/3 (7499 sayılı Kanun ile değişik) — ad ve metinde aranacak ifade
SEBEPLER = {
    "5/1": ("Açık rıza", "acik riza"),
    "5/2-a": ("Kanunlarda açıkça öngörülmesi", "kanunlarda acikca ongorul"),
    "5/2-b": ("Fiilî imkânsızlık hâlinde hayatın veya beden bütünlüğünün korunması", "beden butunlug"),
    "5/2-c": ("Bir sözleşmenin kurulması veya ifasıyla doğrudan doğruya ilgili olması", "sozlesmenin kurulmasi veya ifasi"),
    "5/2-ç": ("Veri sorumlusunun hukuki yükümlülüğünü yerine getirebilmesi için zorunlu olması", "hukuki yukumlulu"),
    "5/2-d": ("İlgili kişinin kendisi tarafından alenileştirilmiş olması", "alenilestir"),
    "5/2-e": ("Bir hakkın tesisi, kullanılması veya korunması için zorunlu olması", "hakkin tesisi"),
    "5/2-f": ("İlgili kişinin temel hak ve özgürlüklerine zarar vermemek kaydıyla veri sorumlusunun meşru menfaatleri için zorunlu olması", "mesru menfaat"),
    "6/3-a": ("Açık rıza", "acik riza"),
    "6/3-b": ("Kanunlarda açıkça öngörülmesi", "kanunlarda acikca ongorul"),
    "6/3-c": ("Fiilî imkânsızlık hâlinde hayatın veya beden bütünlüğünün korunması", "beden butunlug"),
    "6/3-ç": ("İlgili kişinin alenileştirdiği veriler (alenileştirme iradesine uygun olarak)", "alenilestir"),
    "6/3-d": ("Bir hakkın tesisi, kullanılması veya korunması için zorunlu olması", "hakkin tesisi"),
    "6/3-e": ("Kamu sağlığının korunması, tıbbi teşhis, tedavi ve bakım amacıyla sır saklama yükümlülüğü altındaki kişilerce işlenmesi", "sir saklama"),
    "6/3-f": ("İstihdam, iş sağlığı ve güvenliği, sosyal güvenlik alanındaki hukuki yükümlülüklerin yerine getirilmesi", "istihdam"),
    "6/3-g": ("Vakıf, dernek ve diğer kâr amacı gütmeyen kuruluşlarca üyeleri hakkında işlenmesi", "kar amaci gutmeyen"),
}
OZEL = re.compile(r"\b(saglik|tibbi|biyometrik|genetik|ceza mahkumiyet|guvenlik tedbir|din|dini|mezhep|irk|etnik|siyasi|felsefi|inanc|"
                  r"kilik kiyafet|sendika|cinsel|engellilik|kan grubu|parmak izi|yuz tanima)\w*")
YURTDISI = re.compile(r"yeterlilik|standart sozlesme|baglayici sirket|taahhutname|arizi")
GENEL = re.compile(r"\b(vb|vs|vesaire|ve benzeri|ve benzerleri|sinirli olmamak|baslica|baslica|ve diger amaclar|diger amaclar|gibi amaclar|"
                   r"ilgili mevzuat kapsamindaki diger)\b")
RIZA_BEYANI = re.compile(r"onayliyorum|kabul ediyorum|riza veriyorum|onay veriyorum|okudum anladim|muvafakat ediyorum")

SEMA = {
    "type": "object",
    "properties": {
        "giris": S,
        "amac_bolumleri": {"type": "array", "items": {"type": "object", "properties": {"baslik": S, "metin": S, "dayanak": {"type": "array", "items": S}},
                                                      "required": ["baslik", "metin", "dayanak"], "additionalProperties": False}},
        "aktarim": {"type": "array", "items": {"type": "object", "properties": {"metin": S, "dayanak": {"type": "array", "items": S}},
                                               "required": ["metin", "dayanak"], "additionalProperties": False}},
        "toplama_yontemi": S,
    },
    "required": ["giris", "amac_bolumleri", "aktarim", "toplama_yontemi"],
    "additionalProperties": False,
}


def sebep_coz(x: str, ozel: bool) -> str | None:
    """'5/2-c', 'md. 5/2 (ç)', '6/3-f', 'Açık rıza' → kanonik bent kodu."""
    s = str(x or "").lower().replace("i̇", "i")
    m = re.search(r"([56])\s*[/.\-]\s*([123])\s*[\-/.(\s]*([abcçdefg])?(?![a-zçğıöşü])", s)
    if m:
        kod = f"{m.group(1)}/{m.group(2)}" + (f"-{m.group(3)}" if m.group(3) else "")
        return kod if kod in SEBEPLER else None
    if "açık rıza" in s or "acik riza" in katla(s):
        return "6/3-a" if ozel else "5/1"
    return next((k for k, (ad, _) in SEBEPLER.items() if k.startswith("6" if ozel else "5") and katla(ad) == katla(s)), None)


def envanter_oku(yol: Path) -> list[dict]:
    satirlar = [r for r in kayitlar(yol, KOLONLAR, ("kategori", "amac", "sebep")) if r["kategori"] or r["amac"]]
    if not satirlar:
        raise llm.LLMHatasi(f"{yol.name}: envanterde satır yok.")
    for i, r in enumerate(satirlar, 1):
        r["id"] = f"E{i}"
        r = r | {k: str(v) for k, v in r.items()}
        satirlar[i - 1] = r
        r["ozel"] = bool(OZEL.search(katla(r["kategori"] + " " + r["veriler"])))
        r["kod"] = sebep_coz(r["sebep"], r["ozel"])
        r["alicilar"] = [a.strip() for a in re.split(r"\s+/\s+|\|", r["alici"]) if a.strip()]
        yd = katla(r["yurtdisi"])
        r["yd_var"] = bool(yd) and yd not in {"hayir", "yok", "-", "0"}
    return satirlar


def sorumlu_oku(yol: Path) -> dict:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: veri sorumlusu bilgi dosyası bulunamadı.")
    n = {a: "" for a in SORUMLU}
    for s in belge.metin_oku(yol).splitlines():
        a, _, d = s.partition(":")
        k = next((k for k, v in SORUMLU.items() if katla(v) == katla(a)), None)
        if k:
            n[k] = d.strip()
    if not n["unvan"] or not n["adres"]:
        raise llm.LLMHatasi(f"{yol.name}: 'Unvan:' ve 'Adres:' satırları gerekli (veri sorumlusunun kimliği, KVKK md. 10).")
    return n


def envanter_kontrol(env: list[dict]) -> list[str]:
    """Envanterin kendisindeki sorunlar (metin yazılsa da kalır; envanter düzeltilmeli)."""
    s = []
    for r in env:
        if not r["kod"]:
            s.append(f"{r['id']}: hukuki sebep tanınmadı ('{r['sebep']}'). 5/1, 5/2-a … 5/2-f ya da 6/3-a … 6/3-g yazın")
        elif r["ozel"] and r["kod"].startswith("5"):
            s.append(f"{r['id']}: özel nitelikli veri ({r['kategori']}) md. 5'e dayandırılmış; md. 6/3 bentlerinden biri gerekir")
        elif not r["ozel"] and r["kod"].startswith("6"):
            s.append(f"{r['id']}: özel nitelikli olmayan veri için md. 6/3 yazılmış; md. 5 bendini kontrol edin")
        if r["yd_var"] and not YURTDISI.search(katla(r["yurtdisi"])):
            s.append(f"{r['id']}: yurt dışı aktarımın md. 9 mekanizması belirsiz ('{r['yurtdisi']}'): yeterlilik kararı, standart sözleşme, "
                     "bağlayıcı şirket kuralları, taahhütname ya da arızi aktarım hâli yazın")
        if not r["alicilar"] and r["aktarim_amaci"]:
            s.append(f"{r['id']}: aktarım amacı var ama alıcı grubu boş")
    return s


def denetle(t: dict, env: list[dict]) -> list[str]:
    sorun = []
    e = {r["id"]: r for r in env}
    kodlar = {r["kod"] for r in env if r["kod"]}
    kullanilan = set()
    for i, b in enumerate(t["amac_bolumleri"], 1):
        ad = f"Bölüm {i} ({b['baslik'][:40]})"
        gec = [d for d in b["dayanak"] if d not in e]
        if not b["dayanak"] or gec:
            sorun.append(f"{ad}: dayanak yok ya da geçersiz {gec}")
        kullanilan |= set(b["dayanak"])
        km = katla(b["metin"])
        for d in b["dayanak"]:
            if d in e and e[d]["kod"] and SEBEPLER[e[d]["kod"]][1] not in km:
                sorun.append(f"{ad}: {d} satırının hukuki sebebi ({SEBEPLER[e[d]['kod']][0]}) metinde yazılmamış")
            if d in e and not metinde_var(e[d]["amac"], b["metin"]) and not _ortak_kelime(e[d]["amac"], b["metin"]):
                sorun.append(f"{ad}: {d} satırının işleme amacı ('{e[d]['amac'][:50]}') metinde yok")
    for r in env:
        if r["id"] not in kullanilan:
            sorun.append(f"{r['id']}: envanter satırı ({r['kategori']}, {r['amac'][:40]}) hiçbir bölümde ele alınmamış")
    tum = katla(" ".join([t["giris"], t["toplama_yontemi"]] + [b["metin"] for b in t["amac_bolumleri"]] + [a["metin"] for a in t["aktarim"]]))
    for kod, (ad, ifade) in SEBEPLER.items():
        if ifade in tum and not any(SEBEPLER[k][1] == ifade for k in kodlar):
            sorun.append(f"Metin: envanterde olmayan hukuki sebep yazılmış ({ad})")
    aktarim = katla(" ".join(a["metin"] for a in t["aktarim"]))
    for r in env:
        for a in r["alicilar"]:
            if not metinde_var(a, aktarim):
                sorun.append(f"Aktarım: {r['id']} alıcı grubu '{a}' yazılmamış")
        if r["yd_var"] and ("yurt disi" not in aktarim or not YURTDISI.search(aktarim)):
            sorun.append(f"Aktarım: {r['id']} yurt dışına aktarılıyor; yurt dışı aktarım ve mekanizması ({r['yurtdisi']}) yazılmamış")
    for y in dict.fromkeys(r["yontem"] for r in env if r["yontem"]):
        if not metinde_var(y, t["toplama_yontemi"]):
            sorun.append(f"Toplama yöntemi: '{y}' yazılmamış")
    if GENEL.search(tum):
        sorun.append(f"Metin: genel / belirsiz ifade ('{GENEL.search(tum).group(0)}'); amaç ve alıcıları tek tek ve açıkça yazın")
    if RIZA_BEYANI.search(tum):
        sorun.append("Metin: aydınlatma metnine açık rıza beyanı eklenmiş; açık rıza ayrı bir metinle alınmalıdır")
    return list(dict.fromkeys(sorun))


def _ortak_kelime(amac: str, metin: str) -> bool:
    """Amacın anlamlı kelimelerinin en az yarısı metinde geçiyor mu (eş anlamlı yeniden yazıma izin)."""
    k = [w[:6] for w in katla(amac).split() if len(w) > 3 and w not in {"veya", "ilgili", "olarak"}]
    m = katla(metin)
    return bool(k) and sum(w in m for w in k) * 2 >= len(k)


def uret(env: list[dict], s: dict, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    satir = "\n".join(f"{r['id']} | {r['kategori']} | {r['veriler']} | amaç: {r['amac']} | hukuki sebep: "
                      f"{SEBEPLER[r['kod']][0] + ' (KVKK md. ' + r['kod'] + ')' if r['kod'] else '(tanınmadı: ' + r['sebep'] + ')'} | "
                      f"yöntem: {r['yontem']} | alıcı: {', '.join(r['alicilar']) or '-'} | aktarım amacı: {r['aktarim_amaci'] or '-'} | "
                      f"yurt dışı: {r['yurtdisi'] or 'yok'}" for r in env)
    mesaj = (f"<veri_sorumlusu>{s['unvan']}</veri_sorumlusu>\n<ilgili_kisi_grubu>{s['grup'] or '(belirtilmemiş)'}</ilgili_kisi_grubu>\n"
             f"<envanter>\n{satir}\n</envanter>" + duzeltme_blogu(duzeltme, onceki))
    return llm.json_iste(sistem, mesaj, SEMA)


HAKLAR = [
    "kişisel verilerinizin işlenip işlenmediğini öğrenme,", "işlenmişse buna ilişkin bilgi talep etme,",
    "işlenme amacını ve bunların amacına uygun kullanılıp kullanılmadığını öğrenme,", "yurt içinde veya yurt dışında aktarıldığı üçüncü kişileri bilme,",
    "eksik veya yanlış işlenmiş olması hâlinde bunların düzeltilmesini isteme,",
    "KVKK md. 7'de öngörülen şartlar çerçevesinde silinmesini veya yok edilmesini isteme,",
    "düzeltme, silme ve yok etme işlemlerinin kişisel verilerin aktarıldığı üçüncü kişilere bildirilmesini isteme,",
    "işlenen verilerin münhasıran otomatik sistemler vasıtasıyla analiz edilmesi suretiyle aleyhinize bir sonucun ortaya çıkmasına itiraz etme,",
    "kişisel verilerin kanuna aykırı olarak işlenmesi sebebiyle zarara uğramanız hâlinde zararın giderilmesini talep etme.",
]


def metin_olustur(t: dict, s: dict) -> str:
    kim = f"{s['unvan']}" + (f" (MERSİS No: {s['mersis']})" if s["mersis"] else "") + f", Adres: {s['adres']}"
    y = [f"{s['unvan'].upper()} KİŞİSEL VERİLERİN İŞLENMESİNE İLİŞKİN AYDINLATMA METNİ", ""]
    if s["grup"]:
        y += [f"İlgili kişi grubu: {s['grup']}", ""]
    y += ["1. Veri Sorumlusu", f"6698 sayılı Kişisel Verilerin Korunması Kanunu (\"KVKK\") uyarınca veri sorumlusu {kim}'dir.", "", t["giris"], "",
          "2. İşlenen Kişisel Veriler, İşleme Amaçları ve Hukuki Sebepler"]
    for b in t["amac_bolumleri"]:
        y += ["", b["baslik"], b["metin"]]
    y += ["", "3. Kişisel Verilerin Aktarılması"] + [a["metin"] for a in t["aktarim"]]
    y += ["", "4. Kişisel Verilerin Toplanma Yöntemi", t["toplama_yontemi"]]
    y += ["", "5. KVKK md. 11 Kapsamındaki Haklarınız", "KVKK'nın 11. maddesi uyarınca veri sorumlusuna başvurarak:"]
    y += [f"  {h}" for h in HAKLAR]
    kanal = [f"yazılı olarak {s['adres']} adresine"]
    if s["kep"]:
        kanal.append(f"kayıtlı elektronik posta (KEP) ile {s['kep']} adresine")
    kanal.append("güvenli elektronik imza ya da mobil imza ile")
    if s["eposta"]:
        kanal.append(f"daha önce tarafımıza bildirdiğiniz ve sistemimizde kayıtlı e-posta adresinizden {s['eposta']} adresine")
    y += ["haklarına sahipsiniz.", "", "6. Başvuru",
          "Taleplerinizi Veri Sorumlusuna Başvuru Usul ve Esasları Hakkında Tebliğ'e uygun olarak; " + "; ".join(kanal) + " iletebilirsiniz. "
          "Başvurunuz, niteliğine göre en kısa sürede ve en geç otuz gün içinde ücretsiz olarak sonuçlandırılır; işlemin ayrıca bir maliyet "
          "gerektirmesi hâlinde Kişisel Verileri Koruma Kurulunca belirlenen tarifedeki ücret alınabilir.",
          "", f"Yürürlük tarihi: {date.today():%d.%m.%Y}"]
    return "\n".join(y)


def rapor_yaz(cikti: Path, t, env, s, env_sorun, kalan, tur, metin) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Metin"
    baslik(ws, ["Bölüm", "Metin", "Dayanak"], (30, 100, 14))
    ws.append(["Giriş", t["giris"], ""])
    for b in t["amac_bolumleri"]:
        ws.append([b["baslik"], b["metin"], ", ".join(b["dayanak"])])
    for a in t["aktarim"]:
        ws.append(["Aktarım", a["metin"], ", ".join(a["dayanak"])])
    ws.append(["Toplama yöntemi", t["toplama_yontemi"], ""])
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 2).fill = AI_DOLGU
    son_not(ws, "Önemli: Taslaktır. Envanter, hukuki sebepler ve aktarımlar hukuk birimi / avukat tarafından doğrulanmadan yayımlanmamalıdır.", kalan, tur)
    hizala(ws)
    w = wb.create_sheet("Envanter Kontrol")
    baslik(w, ["No", "Kategori", "Veriler", "Amaç", "Hukuki Sebep (girilen)", "Hukuki Sebep (KVKK)", "Özel Nitelikli", "Alıcılar", "Yurt Dışı", "Kontrol"],
           (6, 16, 30, 40, 16, 50, 10, 30, 16, 50))
    for r in env:
        sorun = [x.split(": ", 1)[1] for x in env_sorun if x.startswith(r["id"] + ":")]
        w.append([r["id"], r["kategori"], r["veriler"], r["amac"], r["sebep"], f"{r['kod']}: {SEBEPLER[r['kod']][0]}" if r["kod"] else "?",
                  "Evet" if r["ozel"] else "Hayır", ", ".join(r["alicilar"]), r["yurtdisi"], "; ".join(sorun) or "Uygun"])
        boya(w, (10,), KIRMIZI if sorun else YESIL)
    hizala(w)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    cikti.with_suffix(".txt").write_text(metin, encoding="utf-8")


def calistir(envanter: Path, sorumlu: Path, cikti: Path, tur: int = 2, evet: bool = False) -> dict:
    env = envanter_oku(envanter)
    s = sorumlu_oku(sorumlu)
    env_sorun = envanter_kontrol(env)
    print(f"[OK] {s['unvan']} · {len(env)} envanter satırı")
    for x in env_sorun:
        print(f"[!] Envanter: {x}")
    llm.onay_al(f"Veri işleme envanteri ({len(env)} satır) ve veri sorumlusu unvanı gönderilecek; kişisel veri içermez.", evet)
    t, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(env, s, d, o), lambda x: denetle(x, env), tur)
    metin = metin_olustur(t, s)
    rapor_yaz(cikti, t, env, s, env_sorun, kalan, yapilan, metin)
    return {"envanter": env, "envanter_sorun": env_sorun, "taslak": t, "kalan": kalan, "tur": yapilan, "metin": metin}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    p = argparse.ArgumentParser(description="Veri işleme envanterinden KVKK md. 10'a uygun aydınlatma metni taslağı hazırlar.")
    v = BURASI / "ornek_veri"
    p.add_argument("--envanter", type=Path, default=v / "veri_envanteri.csv",
                   help="Envanter (CSV/Excel): Veri Kategorisi, Veriler, İşleme Amacı, Hukuki Sebep, Toplama Yöntemi, Alıcı Grubu, Aktarım Amacı, Yurt Dışı Aktarım")
    p.add_argument("--sorumlu", type=Path, default=v / "veri_sorumlusu.txt", help="Veri sorumlusu (.txt/.docx): Unvan, MERSİS, Adres, KEP, E-posta, İlgili Kişi Grubu")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "aydinlatma_metni.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.envanter, a.sorumlu, a.cikti, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(s['taslak']['amac_bolumleri'])} amaç bölümü · {len(s['envanter_sorun'])} envanter uyarısı")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Metin: {a.cikti.with_suffix('.txt').resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
