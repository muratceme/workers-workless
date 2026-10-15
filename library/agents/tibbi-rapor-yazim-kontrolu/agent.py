"""
Tıbbi Rapor Yazım Kontrolü — Workers / Workless AI Agent
Tıbbi Sekreterlik › Tıbbi Sekreter

1. Kod hasta kaydını (ad, TC kimlik no, doğum tarihi, cinsiyet, protokol no, rapor tarihi, rapor türü) ve hekimin dikte
   ettiği rapor metnini okur. Hasta adı {HASTA} olur; TC kimlik no, telefon gibi bilgiler maskelenir. Hasta kaydı modele
   gönderilmez.
2. Kod tutarlılık kontrollerini kendisi yapar:
   - Metindeki yaş ("45 yaşında") doğum tarihine göre doğru mu; cinsiyet ifadeleri (kadın / erkek hasta, gebe, prostat …)
     kayıtla çelişiyor mu
   - Metindeki tarihler rapor tarihinden sonra ya da doğum tarihinden önce mi
   - Metindeki TC kimlik no ve protokol no kayıtla aynı mı
   - Rapor türü tek taraflıysa (ör. "Sağ diz") metinde karşı taraf geçiyor mu
3. Model yazım, tıbbi terim, taraf tutarlılığı, birim ve kısaltma hatalarını bulur: her biri için metinden aynen alıntı,
   öneri ve açıklama. Model metni yeniden yazmaz.
4. Kod modelin bulgularını denetler (en fazla --tur kez geri gönderilir):
   - Alıntı metinde aynen geçiyor mu; öneri alıntıdan farklı mı
   - Öneri bir sayıyı değiştiriyorsa "hekime sor" işaretli mi (ölçü, doz, tarih hekim onayı olmadan düzeltilemez)
   - Yazım önerisi alıntıya yeterince benziyor mu (anlam değiştiren "düzeltme" yapılmasın)
5. Kod "hekime sor" olmayan önerileri metne uygular ve önerilen metni ayrıca kaydeder. Asıl rapor değişmez.

Kullanım:
    python agent.py                                              # örnek (kurgusal) diz MR raporu
    python agent.py --hasta hasta.txt --rapor dikte.docx
"""
from __future__ import annotations

import argparse
import csv
import difflib
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
HASTA = {"ad": "Hasta Adı", "tc": "TC Kimlik No", "dogum": "Doğum Tarihi", "cinsiyet": "Cinsiyet", "protokol": "Protokol No",
         "tarih": "Rapor Tarihi", "tur": "Rapor Türü"}
TURLER = ["Yazım", "Tıbbi terim", "Taraf tutarlılığı", "Birim / ölçü", "Kısaltma", "Noktalama", "Diğer"]
KADIN = re.compile(r"\b(kadın hasta|bayan hasta|gebe\w*|hamile\w*|uterus\w*|over(ler)?\b|endometri\w*|serviks\w*|vajen\w*|menopoz\w*)")
ERKEK = re.compile(r"\b(erkek hasta|bay hasta|prostat\w*|testis\w*|skrotum\w*|skrotal)")
TARAF = {"sağ": "sol", "sol": "sağ"}


def kucuk(metin: str) -> str:
    """Türkçe küçük harf; uzunluk korunur (konumlar asıl metinle aynı kalır)."""
    return metin.replace("I", "ı").replace("İ", "i").lower()

SEMA = {
    "type": "object",
    "properties": {
        "bulgular": {"type": "array", "items": {"type": "object", "properties": {
            "tur": {"type": "string", "enum": TURLER}, "alinti": S, "oneri": S, "aciklama": S, "hekime_sor": {"type": "boolean"}},
            "required": ["tur", "alinti", "oneri", "aciklama", "hekime_sor"], "additionalProperties": False}},
    },
    "required": ["bulgular"],
    "additionalProperties": False,
}


def hasta_oku(yol: Path) -> dict:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: hasta kaydı bulunamadı.")
    h = {a: "" for a in HASTA}
    for s in belge.metin_oku(yol).splitlines():
        a, _, d = s.partition(":")
        k = next((k for k, v in HASTA.items() if katla(v) == katla(a)), None)
        if k:
            h[k] = d.strip()
    h["dogum_v"], h["tarih_v"] = tarih_coz(h["dogum"]), tarih_coz(h["tarih"]) or date.today()
    if not h["dogum_v"] or katla(h["cinsiyet"]) not in {"kadin", "erkek"}:
        raise llm.LLMHatasi(f"{yol.name}: 'Doğum Tarihi: GG.AA.YYYY' ve 'Cinsiyet: Kadın/Erkek' satırları gerekli.")
    return h


def yas(dogum: date, gun: date) -> int:
    return gun.year - dogum.year - ((gun.month, gun.day) < (dogum.month, dogum.day))


def gizle(metin: str, h: dict) -> str:
    if h["ad"]:
        metin = re.sub(re.escape(h["ad"]), "{HASTA}", metin, flags=re.I)
        for p in h["ad"].split():
            if len(p) > 2:
                metin = re.sub(rf"\b{re.escape(p)}\b", "{HASTA}", metin)
    metin = re.sub(r"(protokol\w*\s*(?:no)?\s*[:.]?\s*)[\w-]+", r"\1[PROTOKOL]", metin, flags=re.I)
    if h["protokol"]:
        metin = metin.replace(h["protokol"], "[PROTOKOL]")
    return llm.maskele(metin)


def cumle_bul(metin: str, i: int) -> str:
    bas = max(metin.rfind(".", 0, i), metin.rfind("\n", 0, i)) + 1
    son = min([x for x in (metin.find(".", i), metin.find("\n", i)) if x >= 0] or [len(metin)])
    return metin[bas:son + 1].strip()


def kod_kontrolleri(metin: str, h: dict) -> list[dict]:
    out = []

    def ekle(tur, alinti, aciklama):
        out.append({"tur": tur, "alinti": alinti, "oneri": "", "aciklama": aciklama, "hekime_sor": True, "kaynak": "Kod"})

    beklenen = yas(h["dogum_v"], h["tarih_v"])
    for m in re.finditer(r"(\d{1,3})\s*yaş", metin):
        if int(m.group(1)) != beklenen:
            ekle("Hasta bilgisi", m.group(0), f"Doğum tarihine ({h['dogum']}) göre hasta rapor tarihinde {beklenen} yaşında")
    k = kucuk(metin)
    desen = KADIN if katla(h["cinsiyet"]) == "erkek" else ERKEK
    for m in desen.finditer(k):
        ifade = metin[m.start():m.end()]
        ekle("Hasta bilgisi", ifade, f"Kayıtta cinsiyet {h['cinsiyet']}; '{ifade}' ifadesi kayıtla çelişiyor")
    for m in re.finditer(r"\b\d{2}\.\d{2}\.\d{4}\b", metin):
        t = tarih_coz(m.group(0))
        if t and t > h["tarih_v"]:
            ekle("Tarih", m.group(0), f"Rapor tarihinden ({h['tarih']}) sonra")
        elif t and t < h["dogum_v"]:
            ekle("Tarih", m.group(0), f"Doğum tarihinden ({h['dogum']}) önce")
    for m in re.finditer(r"\b[1-9]\d{10}\b", metin):
        if h["tc"] and m.group(0) != h["tc"]:
            ekle("Hasta bilgisi", m.group(0), "TC kimlik no hasta kaydıyla aynı değil")
    for m in re.finditer(r"protokol\w*\s*(?:no)?\s*[:.]?\s*([\w-]+)", metin, flags=re.I):
        if h["protokol"] and m.group(1) != h["protokol"]:
            ekle("Hasta bilgisi", m.group(1), f"Protokol no hasta kaydıyla ({h['protokol']}) aynı değil")
    taraf = next((t for t in TARAF if re.search(rf"\b{t}\b", kucuk(h["tur"]))), None)
    if taraf:
        for m in re.finditer(rf"\b{TARAF[taraf]}\b", k):
            ekle("Taraf tutarlılığı", cumle_bul(metin, m.start()), f"Rapor türü '{h['tur']}'; metinde karşı taraf ({TARAF[taraf]}) geçiyor")
    return out


def denetle(t: dict, metin: str) -> list[str]:
    sorun = []
    gorulen = set()
    for i, b in enumerate(t["bulgular"], 1):
        ad = f"Bulgu {i} ('{b['alinti'][:40]}')"
        if not b["alinti"] or b["alinti"] not in metin:
            sorun.append(f"{ad}: alıntı metinde aynen geçmiyor (büyük / küçük harf ve noktalama dahil aynen alın)")
            continue
        if b["alinti"] in gorulen:
            sorun.append(f"{ad}: aynı alıntı iki kez yazılmış")
        gorulen.add(b["alinti"])
        if b["hekime_sor"]:
            continue
        if not b["oneri"] or b["oneri"] == b["alinti"]:
            sorun.append(f"{ad}: öneri boş ya da alıntıyla aynı; düzeltilemiyorsa hekime_sor işaretleyin")
        elif sayilar(b["oneri"]) != sayilar(b["alinti"]):
            sorun.append(f"{ad}: öneri bir sayıyı değiştiriyor; ölçü / doz / tarih hekim onayı olmadan düzeltilemez, hekime_sor işaretleyin")
        elif b["tur"] in ("Yazım", "Noktalama") and difflib.SequenceMatcher(None, b["alinti"].lower(), b["oneri"].lower()).ratio() < 0.6:
            sorun.append(f"{ad}: yazım önerisi alıntıdan çok farklı; anlam değişiyorsa hekime_sor işaretleyin")
    return sorun


def uret(metin: str, h: dict, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    mesaj = f"<rapor_turu>{h['tur'] or '-'}</rapor_turu>\n<rapor>\n{metin}\n</rapor>" + duzeltme_blogu(duzeltme, onceki)
    return llm.json_iste(sistem, mesaj, SEMA)


def onerilen_metin(asil: str, bulgular: list[dict]) -> tuple[str, int]:
    metin, n = asil, 0
    for b in bulgular:
        if not b["hekime_sor"] and b["oneri"] and b["alinti"] in metin:
            metin = metin.replace(b["alinti"], b["oneri"], 1)
            n += 1
    return metin, n


def birlestir(kod: list[dict], model: list[dict]) -> list[dict]:
    out = list(kod)
    for b in model:
        if any(b["alinti"] in k["alinti"] or k["alinti"] in b["alinti"] for k in kod if k["tur"] == b["tur"]):
            continue
        out.append(b | {"kaynak": "Model"})
    return out


def rapor_yaz(cikti: Path, bulgular, asil, oneri, h, kalan, tur) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Bulgular"
    baslik(ws, ["No", "Kaynak", "Tür", "Alıntı", "Öneri", "Açıklama", "Hekime Sor"], (5, 8, 16, 50, 40, 50, 10))
    for i, b in enumerate(bulgular, 1):
        ws.append([i, b["kaynak"], b["tur"], b["alinti"], b["oneri"], b["aciklama"], "Evet" if b["hekime_sor"] else ""])
        boya(ws, (7,), KIRMIZI if b["hekime_sor"] else SARI)
        if b["kaynak"] == "Model":
            boya(ws, (5, 6), AI_DOLGU)
    son_not(ws, f"Rapor türü: {h['tur']} · Önemli: Öneriler hekim onayı olmadan rapora işlenmez. 'Hekime Sor' işaretli bulgular hekimle netleştirilmelidir.",
            kalan, tur)
    hizala(ws)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    cikti.with_suffix(".txt").write_text("ÖNERİLEN METİN (hekim onayına sunulacaktır; asıl rapor değiştirilmemiştir)\n\n" + oneri, encoding="utf-8")


def calistir(hasta_yol: Path, rapor_yol: Path, cikti: Path, tur: int = 2, evet: bool = False) -> dict:
    h = hasta_oku(hasta_yol)
    if not rapor_yol.exists():
        raise llm.LLMHatasi(f"{rapor_yol}: rapor metni bulunamadı.")
    asil = belge.metin_oku(rapor_yol)
    kod = kod_kontrolleri(asil, h)
    gonderilen = gizle(asil, h)
    print(f"[OK] {h['tur'] or 'Rapor'} · {len(asil.split())} kelime · kodun bulduğu tutarsızlık: {len(kod)}")
    llm.onay_al("Rapor metni gönderilecek (hasta adı, TC kimlik no ve protokol no maskelenir; hasta kaydı gönderilmez). Sağlık verisi "
                "içerdiğinden kurumunuzun izni olmadan çalıştırmayın.", evet)
    t, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(gonderilen, h, d, o), lambda x: denetle(x, gonderilen), tur)
    model = [b for b in t["bulgular"] if b["alinti"] in asil or b["alinti"] in gonderilen]
    bulgular = birlestir(kod, model)
    oneri, n = onerilen_metin(asil, [b for b in model if b["alinti"] in asil])
    rapor_yaz(cikti, bulgular, asil, oneri, h, kalan, yapilan)
    return {"hasta": h, "bulgular": bulgular, "kalan": kalan, "tur": yapilan, "oneri": oneri, "uygulanan": n, "gonderilen": gonderilen}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    p = argparse.ArgumentParser(description="Hekimin dikte ettiği rapor metnini yazım, tarih ve hasta bilgisi tutarlılığı açısından kontrol eder.")
    v = BURASI / "ornek_veri"
    p.add_argument("--hasta", type=Path, default=v / "hasta.txt",
                   help="Hasta kaydı (.txt/.docx): Hasta Adı, TC Kimlik No, Doğum Tarihi, Cinsiyet, Protokol No, Rapor Tarihi, Rapor Türü")
    p.add_argument("--rapor", type=Path, default=v / "rapor.txt", help="Rapor metni (.txt/.docx/.pdf)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "rapor_kontrol.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.hasta, a.rapor, a.cikti, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(s['bulgular'])} bulgu · {sum(b['hekime_sor'] for b in s['bulgular'])} hekime sorulacak · {s['uygulanan']} öneri metne uygulandı")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
