"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri" / "fit_yorumlari.txt"
MESAJLAR = []


def a(model, bolge, birim, satir, olcu=None, onc="Normal"):
    return {"model": model, "bolge": bolge, "yorum": "yorum", "aksiyon": "aksiyon", "birim": birim, "oncelik": onc, "olcu": olcu, "kaynak_satirlar": satir}


def o(pom, olculen, istenen, degisim, birim):
    return {"pom": pom, "olculen": olculen, "istenen": istenen, "degisim": degisim, "birim": birim}


def taslak(dogru=True) -> dict:
    t = {"kararlar": [{"model": "WW-2471", "tanim": "Kadın örme elbise", "numune": "Fit", "beden": "M", "karar": "Yorumlarla onaylandı", "kaynak_satirlar": [2, 3]},
                      {"model": "WW-2488", "tanim": "Erkek polo", "numune": "Belirtilmemiş", "beden": "", "karar": "Reddedildi / yeniden numune", "kaynak_satirlar": [10]}],
         "aksiyonlar": [a("WW-2471", "Göğüs eni", "Kalıp", [4], o("Göğüs eni", 52, 50, -2, "cm")),
                        a("WW-2471", "Kol boyu", "Kalıp", [5], o("Kol boyu", None, None, 0.75, "inç")),
                        a("WW-2471", "Yaka", "Modelhane", [6]),
                        a("WW-2471", "Etek ucu", "Kumaş / aksesuar", [7]),
                        a("WW-2471", "Renk", "Kumaş / aksesuar", [8]),
                        a("WW-2471", "Göğüs baskısı", "Baskı / nakış", [9], o("Baskı yeri", None, None, 2, "cm")),
                        a("WW-2488", "Kumaş gramajı", "Kumaş / aksesuar", [10], o("Gramaj", 165, 180, None, "gsm"), "Yüksek")],
         "terminler": [{"model": "WW-2471", "istenen": "2 adet size set numunesi", "tarih": "25.10.2026", "kaynak_satirlar": [11]}]}
    if not dogru:
        t["aksiyonlar"][1]["olcu"] = o("Kol boyu", None, None, 1.9, "cm")
        del t["aksiyonlar"][3]
        t["terminler"][0]["tarih"] = "26.10.2026"
    return t


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return taslak("<duzeltme>" in kullanici)


class KodTesti(unittest.TestCase):
    def test_yardimcilar(self):
        self.assertEqual(agent.olculer('increase by 3/4", hem 1 1/2" and 52 cm'), {"0.75", "1.5", "52"})
        self.assertEqual(agent.olculer("received 08/10/2026"), agent.sayilar("received 08/10/2026"))
        self.assertEqual((agent.cm(0.75, "inç"), agent.cm(2, "cm"), agent.cm(165, "gsm")), (1.91, 2, None))
        l = agent.satirlar(ORNEK.read_text(encoding="utf-8"))
        self.assertEqual([i for i, s in enumerate(l, 1) if agent.selam_mi(s, i, len(l))], [1, 12, 13])
        self.assertFalse(agent.selam_mi("Rejected.", 5, 13))

    def test_denetim(self):
        l = agent.satirlar(ORNEK.read_text(encoding="utf-8"))
        self.assertEqual(agent.denetle(taslak(), l), [])
        s = agent.denetle(taslak(False), l)
        self.assertEqual(len(s), 3)
        self.assertTrue(any("1.9" in x for x in s) and any(x.startswith("Satır 7") for x in s) and any("26.10.2026" in x for x in s))
        t = taslak()
        t["aksiyonlar"][0]["kaynak_satirlar"] = [4, 40]
        t["kararlar"][1]["model"] = "WW-2490"
        self.assertEqual(len(agent.denetle(t, l)), 2)


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "f.xlsx"
            s = agent.calistir(ORNEK, cikti, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            self.assertIn('5: Sleeve length is short, please increase by 3/4".', MESAJLAR[0])
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Kararlar ve Terminler", "Aksiyon Listesi", "E-posta"])
            ws = wb["Aksiyon Listesi"]
            kol = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 3).value == "Kol boyu")
            self.assertEqual((ws.cell(kol, 11).value, ws.cell(kol, 13).value), (0.75, 1.91))
            md = s["md"].read_text(encoding="utf-8")
            self.assertIn("(= 1.91 cm)", md)
            self.assertIn("## Kalıp", md)

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(ORNEK, Path(tmp) / "f.xlsx", tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 3)
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--girdi", str(ORNEK.with_name("yok.txt")), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
