"""
Müşteri Yorumları (Fit Comments) Özeti — Gerisi Bende AI Agent
Müşteri Temsilciliği (Merchandising) › Development Merchandiser

1. Kod müşterinin fit / numune yorum e-postasını satır satır numaralar; selamlama ve imza satırlarını ayırır, satırlardaki
   ölçüleri (kesirli inç dahil: 3/4", 1 1/2") bulur.
2. Model yorumları modelhane ve üretim için madde madde aksiyon listesine çevirir: model (style) no, numune türü, bölge,
   yorumun Türkçesi, yapılacak iş, sorumlu birim (modelhane, kalıp, kumaş, baskı / nakış, kalite, üretim, müşteri
   temsilcisi), ölçü değişikliği (POM, ölçülen, istenen, değişim, birim) ve dayandığı satır. Model başına genel kararı
   (onay, yorumlu onay, ret) ve istenen terminleri de çıkarır.
3. Kod denetler: her madde geçerli satıra dayanıyor mu, model no e-postada geçiyor mu, ölçü değerleri dayandığı satırda
   yazıyor mu, terminin tarihi satırda var mı, selamlama / imza dışında kullanılmayan satır kalmış mı. Sorun varsa
   modele geri gönderilir (en fazla --tur kez). İnç değerler kodla cm'ye çevrilir (1 inç = 2,54 cm).

Kullanım:
    python agent.py                                        # örnek fit yorumu e-postasıyla dener
    python agent.py --girdi fit_comments.docx --tur 3
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
BIRIMLER = ["Modelhane", "Kalıp", "Kumaş / aksesuar", "Baskı / nakış", "Kalite", "Üretim", "Müşteri temsilcisi"]
NUMUNE = ["Proto", "Fit", "Size set", "PP (pre-production)", "TOP (top of production)", "Lab dip / strike off", "Belirtilmemiş"]
KARAR = ["Onaylandı", "Yorumlarla onaylandı", "Reddedildi / yeniden numune", "Karar belirtilmemiş"]
SELAM = re.compile(r"^(hi|hello|dear|merhaba|sayın|best regards|kind regards|regards|thanks|thank you|cheers|saygılar|iyi çalışmalar|teşekkürler)\b", re.I)
N = {"type": ["number", "null"]}
SEMA = {
    "type": "object",
    "properties": {
        "kararlar": {"type": "array", "items": {"type": "object", "properties": {
            "model": S, "tanim": S, "numune": {"type": "string", "enum": NUMUNE}, "beden": S, "karar": {"type": "string", "enum": KARAR}, "kaynak_satirlar": {"type": "array", "items": {"type": "integer"}}},
            "required": ["model", "tanim", "numune", "beden", "karar", "kaynak_satirlar"], "additionalProperties": False}},
        "aksiyonlar": {"type": "array", "items": {"type": "object", "properties": {
            "model": S, "bolge": S, "yorum": S, "aksiyon": S, "birim": {"type": "string", "enum": BIRIMLER},
            "oncelik": {"type": "string", "enum": ["Yüksek", "Normal"]},
            "olcu": {"type": ["object", "null"], "properties": {"pom": S, "olculen": N, "istenen": N, "degisim": N, "birim": {"type": "string", "enum": ["cm", "inç", "gsm", "diğer"]}},
                     "required": ["pom", "olculen", "istenen", "degisim", "birim"], "additionalProperties": False},
            "kaynak_satirlar": {"type": "array", "items": {"type": "integer"}}},
            "required": ["model", "bolge", "yorum", "aksiyon", "birim", "oncelik", "olcu", "kaynak_satirlar"], "additionalProperties": False}},
        "terminler": {"type": "array", "items": {"type": "object", "properties": {"model": S, "istenen": S, "tarih": S, "kaynak_satirlar": {"type": "array", "items": {"type": "integer"}}},
                                                 "required": ["model", "istenen", "tarih", "kaynak_satirlar"], "additionalProperties": False}},
    },
    "required": ["kararlar", "aksiyonlar", "terminler"],
    "additionalProperties": False,
}
KESIR = re.compile(r"(?:(\d+)\s+)?(\d+)\s*/\s*(2|3|4|8|16|32)(?![\d/])")


def satirlar(metin: str) -> list[str]:
    return [s.strip() for s in metin.splitlines() if s.strip()]


def selam_mi(s: str, i: int = 0, n: int = 0) -> bool:
    """Selamlama / kapanış satırı ya da e-postanın son iki satırındaki kısa imza."""
    return bool(SELAM.match(s.strip())) or (n > 0 and i >= n - 1 and len(s.split()) <= 3 and not re.search(r"\d", s))


def olculer(s: str) -> set[str]:
    """Satırdaki sayılar; kesirler ondalığa çevrilir (3/4 → 0.75, 1 1/2 → 1.5)."""
    out = sayilar(KESIR.sub(" ", s))
    for tam, pay, payda in KESIR.findall(s):
        if int(payda):
            out.add(sayi_yaz((int(tam) if tam else 0) + int(pay) / int(payda)))
    return out


def cm(deger, birim: str):
    if deger is None:
        return None
    return round(deger * 2.54, 2) if birim == "inç" else (deger if birim == "cm" else None)


def denetle(t: dict, l: list[str]) -> list[str]:
    n, metin = len(l), "\n".join(l)
    sorun, kullanilan = [], set()

    def kaynak(ad: str, ks: list[int]) -> list[str]:
        if not ks:
            sorun.append(f"{ad}: kaynak satır numarası yok")
        gec = [k for k in ks if not 1 <= k <= n]
        if gec:
            sorun.append(f"{ad}: e-postada olmayan satır {gec}")
        kullanilan.update(ks)
        return [l[k - 1] for k in ks if 1 <= k <= n]

    for k in t["kararlar"]:
        kaynak(f"Karar {k['model']}", k["kaynak_satirlar"])
        if not metinde_var(k["model"], metin):
            sorun.append(f"Karar: model no '{k['model']}' e-postada geçmiyor")
    for i, a in enumerate(t["aksiyonlar"], 1):
        ad = f"Aksiyon {i} ({a['model']} · {a['bolge']})"
        sat = kaynak(ad, a["kaynak_satirlar"])
        if a["model"] and not metinde_var(a["model"], metin):
            sorun.append(f"{ad}: model no e-postada geçmiyor")
        o = a["olcu"]
        if o:
            degerler = set().union(*(olculer(s) for s in sat)) if sat else set()
            for alan in ("olculen", "istenen", "degisim"):
                if o[alan] is not None and sayi_yaz(abs(o[alan])) not in degerler:
                    sorun.append(f"{ad}: ölçü '{alan}' {sayi_yaz(o[alan])} dayandığı satırda yazmıyor (hesaplamayın, yazılanı aktarın)")
            if all(o[x] is None for x in ("olculen", "istenen", "degisim")):
                sorun.append(f"{ad}: ölçü nesnesi var ama hiçbir değer yok")
    for x in t["terminler"]:
        sat = kaynak(f"Termin {x['model']}", x["kaynak_satirlar"])
        td = tarih_coz(x["tarih"])
        if not td or not any(katla(td.strftime("%d.%m.%Y")) in katla(s) or katla(x["tarih"]) in katla(s) for s in sat):
            sorun.append(f"Termin {x['model']}: tarih '{x['tarih']}' dayandığı satırda yok")
    for i, s in enumerate(l, 1):
        if i not in kullanilan and not selam_mi(s, i, n):
            sorun.append(f"Satır {i} aksiyon listesine girmedi: '{s[:70]}'")
    return sorun


def uret(l: list[str], duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    numarali = "\n".join(f"{i}: {llm.maskele(s)}" for i, s in enumerate(l, 1))
    return llm.json_iste(sistem, f"<eposta>\n{numarali}\n</eposta>" + duzeltme_blogu(duzeltme, onceki), SEMA)


def rapor_yaz(cikti: Path, l, t, kalan, tur) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Aksiyon Listesi"
    baslik(ws, ["No", "Model", "Bölge / Konu", "Müşteri Yorumu (Türkçe)", "Yapılacak", "Sorumlu Birim", "Öncelik", "POM", "Ölçülen", "İstenen", "Değişim",
                "Birim", "Değişim (cm)", "E-posta Satırı", "Durum"], (5, 10, 18, 44, 44, 18, 9, 16, 9, 9, 9, 7, 11, 10, 12))
    for i, a in enumerate(sorted(t["aksiyonlar"], key=lambda a: (BIRIMLER.index(a["birim"]), a["oncelik"] != "Yüksek")), 1):
        o = a["olcu"] or {}
        ws.append([i, a["model"], a["bolge"], a["yorum"], a["aksiyon"], a["birim"], a["oncelik"], o.get("pom", ""), o.get("olculen"), o.get("istenen"),
                   o.get("degisim"), o.get("birim", ""), cm(o.get("degisim"), o.get("birim", "")) if o else None, ", ".join(map(str, a["kaynak_satirlar"])), "Açık"])
        boya(ws, [4, 5, 6], AI_DOLGU)
        if a["oncelik"] == "Yüksek":
            boya(ws, [7], KIRMIZI)
        boya(ws, [15], SARI)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:O{ws.max_row}"
    son_not(ws, "Önemli: Yapay zekâ özetidir. Modelhaneye iletmeden önce müşterinin orijinal e-postasıyla karşılaştırın; ölçü değişikliklerini ölçü tablosuna işleyin.", kalan, tur)
    hizala(ws)
    k = wb.create_sheet("Kararlar ve Terminler", 0)
    baslik(k, ["Model", "Tanım", "Numune", "Beden", "Karar", "E-posta Satırı"], (12, 28, 18, 8, 26, 12))
    for x in t["kararlar"]:
        k.append([x["model"], x["tanim"], x["numune"], x["beden"], x["karar"], ", ".join(map(str, x["kaynak_satirlar"]))])
        boya(k, [5], KIRMIZI if x["karar"].startswith("Reddedildi") else (YESIL if x["karar"] == "Onaylandı" else SARI))
    k.append([])
    baslik(k, ["Model", "İstenen", "Termin", "E-posta Satırı"])
    for x in t["terminler"]:
        k.append([x["model"], x["istenen"], x["tarih"], ", ".join(map(str, x["kaynak_satirlar"]))])
    hizala(k)
    e = wb.create_sheet("E-posta")
    baslik(e, ["Satır", "Metin", "Kullanıldığı Yer"], (6, 100, 30))
    yer = {}
    for x in t["kararlar"]:
        for s in x["kaynak_satirlar"]:
            yer.setdefault(s, []).append("Karar")
    for i, a in enumerate(t["aksiyonlar"], 1):
        for s in a["kaynak_satirlar"]:
            yer.setdefault(s, []).append(f"Aksiyon ({a['bolge']})")
    for x in t["terminler"]:
        for s in x["kaynak_satirlar"]:
            yer.setdefault(s, []).append("Termin")
    for i, s in enumerate(l, 1):
        e.append([i, s, ", ".join(dict.fromkeys(yer.get(i, []))) or ("selamlama / imza" if selam_mi(s, i, len(l)) else "—")])
        if not yer.get(i) and not selam_mi(s, i, len(l)):
            boya(e, [3], KIRMIZI)
    hizala(e)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    md = cikti.with_suffix(".md")
    out = ["# Fit yorumları — aksiyon listesi", ""]
    for x in t["kararlar"]:
        out.append(f"- **{x['model']}** ({x['tanim']}, {x['numune']}, {x['beden']}): {x['karar']}")
    for b in BIRIMLER:
        l2 = [a for a in t["aksiyonlar"] if a["birim"] == b]
        if l2:
            out += ["", f"## {b}"]
            for a in l2:
                o = a["olcu"]
                olc = f" — {o['pom']}: {sayi_yaz(o['degisim'])} {o['birim']}" + (f" (= {cm(o['degisim'], o['birim'])} cm)" if o["birim"] == "inç" and o["degisim"] is not None else "") if o and o["degisim"] is not None else ""
                out.append(f"- [{a['model']}] {a['bolge']}: {a['aksiyon']}{olc}{' **(Yüksek)**' if a['oncelik'] == 'Yüksek' else ''}")
    if t["terminler"]:
        out += ["", "## Terminler"] + [f"- {x['model']}: {x['istenen']} — {x['tarih']}" for x in t["terminler"]]
    md.write_text("\n".join(out) + "\n", encoding="utf-8")
    return md


def calistir(girdi: Path, cikti: Path, tur: int = 2, evet: bool = False) -> dict:
    if not girdi.exists():
        raise belge.BelgeHatasi(f"{girdi}: dosya bulunamadı.")
    l = satirlar(belge.metin_oku(girdi))
    if not l:
        raise belge.BelgeHatasi(f"{girdi.name}: metin bulunamadı.")
    print(f"[OK] {len(l)} satır · selamlama / imza {sum(selam_mi(s, i, len(l)) for i, s in enumerate(l, 1))}")
    llm.onay_al("Fit yorum e-postası (numaralı satırlar; telefon / e-posta maskeli) gönderilecek.", evet)
    t, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(l, d, o), lambda x: denetle(x, l), tur)
    md = rapor_yaz(cikti, l, t, kalan, yapilan)
    return {"satirlar": l, "taslak": t, "kalan": kalan, "tur": yapilan, "md": md}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    p = argparse.ArgumentParser(description="Müşterinin fit ve numune yorum e-postalarını modelhane ve üretim için aksiyon listesine çevirir.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "fit_yorumlari.txt", help="E-posta / yorum metni (.txt, .md, .docx, .pdf)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "fit_aksiyonlari.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.cikti, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    t = s["taslak"]
    print(f"[OK] {len(t['kararlar'])} model kararı · {len(t['aksiyonlar'])} aksiyon · {len(t['terminler'])} termin")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()} (+ {s['md'].name})")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
