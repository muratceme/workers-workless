"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

DOKUM, FORM = agent.ORNEK / "gorusmeler", agent.ORNEK / "kalite_formu.csv"
MESAJLAR = []
G1 = {  # kod → (sonuç, alıntı)
    "K01": ("karsilandi", "Örnek İletişim'e hoş geldiniz, ben Örnek Temsilci 1"),
    "K02": ("karsilandi", "önce kimlik doğrulaması yapmam gerekiyor"),
    "K03": ("karsilandi", "talebinizi özetliyorum efendim"),                       # dökümde yok → inceleme
    "K04": ("karsilandi", "Tamam, teşekkürler."),                                   # müşterinin sözü → inceleme
    "K05": ("karsilandi", "paket bir gün geç kapatılmış"),
    "K06": ("karsilandi", "Talep numaranız 458812"),
    "K07": ("karsilandi", "sizi kısa bir süre bekletebilir miyim"),
    "K08": ("karsilandi", ""),                                                      # ihlal türü: alıntı gerekmez
    "K09": ("karsilandi", "Yardımcı olabileceğim başka bir konu var mı?"),
    "K10": ("karsilandi", ""),
}
G2 = {
    "K01": ("karsilanmadi", ""), "K02": ("karsilanmadi", ""),
    "K03": ("kismen", "Fatura adresinizi hemen değiştiriyorum"),
    "K04": ("karsilanmadi", "Genelde müşteriler bunu yapmadan arıyor."),
    "K05": ("kismen", "bölgenizde çalışma varmış"),
    "K06": ("karsilandi", "Gerek yok, zaten biliniyor."),                           # anahtar ifade yok → inceleme
    "K07": ("karsilanmadi", "Bir saniye."),
    "K08": ("karsilanmadi", "Genelde müşteriler bunu yapmadan arıyor."),
    "K09": ("kismen", "Başka bir şey?"),
    "K10": ("karsilanmadi", "kartınızın üzerindeki 16 haneli numarayı ve şifrenizi alabilir miyim"),
    "K99": ("karsilandi", "x"),
}


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append((sistem, kullanici))
    kaynak = G1 if "hoş geldiniz" in kullanici else G2
    return {"kriterler": [{"kod": k, "sonuc": s, "gerekce": "Gerekçe " + k, "alinti": a} for k, (s, a) in kaynak.items()],
            "guclu_yonler": ["Kimlik doğrulamayı işlem öncesinde yaptı"], "gelisim_alanlari": ["Kapanışta özet"], "geri_bildirim": "Taslak not."}


class KodTesti(unittest.TestCase):
    def test_form_ve_dokum(self):
        k = agent.form_oku(FORM)
        self.assertEqual((len(k), sum(x.puan for x in k)), (10, D(100)))
        self.assertEqual([x.kod for x in k if x.kritik], ["K02", "K10"])
        self.assertEqual([x.kod for x in k if x.ihlal], ["K08", "K10"])
        g = agent.dokum_oku(DOKUM / "gorusme_002.txt")
        self.assertEqual(g.meta["temsilci adi"], "Örnek Temsilci 2")
        self.assertEqual(g.satirlar[0], ("T", "Alo, buyurun."))
        self.assertIn("[TELEFON]", g.satirlar[1][1])
        self.assertEqual(sum(1 for x, _ in g.satirlar if x == "T"), 7)
        self.assertEqual(agent.maskele("kartım 4111 1111 1111 1111"), "kartım [KART]")

    def test_alinti_konumu(self):
        g = agent.dokum_oku(DOKUM / "gorusme_001.txt")
        self.assertEqual(agent.alinti_nerede("talep numaranız 458812", g), "T")
        self.assertEqual(agent.alinti_nerede("Ben o paketi iptal ettirmiştim ama.", g), "M")
        self.assertEqual(agent.alinti_nerede("hiç söylenmemiş bir cümle", g), "")


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_uctan_uca(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "k.xlsx"
            s = agent.calistir(DOKUM, FORM, cikti, evet=True)
            self.assertEqual(len(MESAJLAR), 2)
            self.assertNotIn("0532", MESAJLAR[1][1])
            self.assertIn('kritik="evet" tur="ihlal"', MESAJLAR[0][1])
            g1, g2 = s["gorusmeler"]
            self.assertEqual((g1.sonuclar["K03"]["sonuc"], g1.sonuclar["K04"]["sonuc"]), ("inceleme", "inceleme"))
            self.assertIn("müşterinin sözü", g1.sonuclar["K04"]["not"])
            self.assertEqual(g1.sonuclar["K08"]["sonuc"], "karsilandi")
            self.assertEqual((g1.puan, g1.kritik_hata), (D("100.0"), []))              # 80 / 80 (iki kriter incelemede)
            self.assertEqual(g2.sonuclar["K06"]["sonuc"], "inceleme")                  # anahtar ifade yok
            self.assertEqual(g2.ham_puan, D("19.4"))                                   # (5 + 10 + 2,5) / 90
            self.assertEqual(g2.puan, D(0))
            self.assertEqual(g2.kritik_hata, ["İşlem öncesi kimlik doğrulama", "Gereksiz kişisel veri istememe"])
            self.assertNotIn("K99", g2.sonuclar)
            turler = {(u["tur"], u["kim"]) for u in s["uyarilar"]}
            self.assertTrue({("Kritik hata", "gorusme_002.txt"), ("İnsan incelemesi gerekli", "gorusme_001.txt · K03"),
                             ("İnsan incelemesi gerekli", "gorusme_002.txt · K06")} <= turler)
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Özet", "Kriter Sonuçları", "Geri Bildirim", "Uyarılar"])
            self.assertEqual(wb["Özet"]["D3"].value, 0)

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_kritik_sifirlama_yok_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(DOKUM / "gorusme_002.txt", FORM, Path(tmp) / "k.xlsx", kritik_sifirla=False, evet=True)
            self.assertEqual(s["gorusmeler"][0].puan, D("19.4"))
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(agent.main(["--dokum", str(Path(tmp) / "yok")]), 1)
            bos = Path(tmp) / "bos"
            bos.mkdir()
            self.assertEqual(agent.main(["--evet", "--dokum", str(bos)]), 1)


if __name__ == "__main__":
    unittest.main()
