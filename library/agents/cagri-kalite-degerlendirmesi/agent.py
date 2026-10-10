"""
Çağrı Kalite Değerlendirmesi — Workers / Workless AI Agent
Müşteri Hizmetleri ve Çağrı Merkezi › Müşteri Deneyimi Uzmanı

Görüşme dökümünü kalite formundaki kriterlere göre puanlar ve geri bildirim notu hazırlar:
1. Kod dökümü okur (Temsilci: / Müşteri: satırları), kişisel verileri maskeler (telefon, e-posta, TCKN, IBAN, kart
   numarası) ve konuşma paylarını hesaplar.
2. Model her kriter için sonuç (karşılandı / kısmen / karşılanmadı / uygulanamaz), gerekçe ve dökümden birebir alıntı
   verir; güçlü yönleri, gelişim alanlarını ve geri bildirim notunu yazar.
3. Kod doğrular ve puanlar:
   - "Karşılandı" ve "kısmen" sonuçlarında alıntı zorunludur ve temsilcinin sözlerinde birebir geçmelidir; geçmiyorsa
     kriter "inceleme"ye alınır (puana girmez, insan değerlendirir). Türü "İhlal" olan kriterlerde (yapılmaması gereken
     davranış) tersine, "karşılanmadı" sonucu ihlali gösteren alıntıyla doğrulanır.
   - Formda "Anahtar İfadeler" varsa kod temsilci sözlerinde arar; model "karşılandı" dediği hâlde ifade yoksa kriter
     "inceleme"ye alınır.
   - Puan = Σ ağırlık × (karşılandı 1, kısmen 0,5, karşılanmadı 0) / Σ ağırlık (uygulanamaz ve inceleme hariç) × 100.
   - Kritik kriter karşılanmadıysa "kritik hata": toplam puan 0 olur (--kritik-sifirlama-yok ile yalnız işaretlenir).
Birden çok döküm bir klasörden topluca değerlendirilebilir.

Kullanım:
    python agent.py                                                  # örnek: 2 görüşme, 10 kriterlik form
    python agent.py --dokum gorusmeler/ --form kalite_formu.xlsx
    python agent.py --dokum gorusme_0412.txt --form kalite_formu.csv --kritik-sifirlama-yok
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import llm

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
SONUCLAR = ["karsilandi", "kismen", "karsilanmadi", "uygulanamaz"]
SONUC_ADI = {"karsilandi": "Karşılandı", "kismen": "Kısmen", "karsilanmadi": "Karşılanmadı", "uygulanamaz": "Uygulanamaz", "inceleme": "İnceleme (insan)",
             "yok": "Değerlendirilmedi"}
KATSAYI = {"karsilandi": Decimal(1), "kismen": Decimal("0.5"), "karsilanmadi": Decimal(0)}
TEMSILCI_ADLARI = ("temsilci", "agent", "operator", "musteri temsilcisi", "danisman", "mt")
MUSTERI_ADLARI = ("musteri", "customer", "arayan", "abone", "uye")
_KART = re.compile(r"\b(?:\d[ -]?){13,19}\b")
AZAMI_DOKUM = 60_000

FORM_SUTUNLARI = {"kod": ("kod", "kriter kodu"), "bolum": ("bolum", "kategori"), "kriter": ("kriter",), "aciklama": ("aciklama", "tanim", "beklenen davranis"),
                  "puan": ("puan", "agirlik"), "kritik": ("kritik", "kritik hata"), "ifadeler": ("anahtar ifadeler", "ifadeler"), "tur": ("tur", "kriter turu")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def _normal(s: str) -> str:
    return " ".join(str(s).replace("İ", "i").replace("I", "ı").replace("’", "'").lower().split()).strip(" .,;:!?\"'“”‘’…")


def maskele(s: str) -> str:
    return _KART.sub("[KART]", llm.maskele(s))


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


# ----------------------------------------------------------------------------
# Form ve döküm
# ----------------------------------------------------------------------------

@dataclass
class Kriter:
    kod: str
    bolum: str
    ad: str
    aciklama: str
    puan: Decimal
    kritik: bool
    ifadeler: list[str]
    ihlal: bool = False          # "İhlal" türü: yapılmaması gereken davranış; karşılandı için alıntı aranmaz, karşılanmadı için aranır


def form_oku(yol: Path) -> list[Kriter]:
    satirlar = tablo_oku(yol)
    b = [katla(c) for c in satirlar[0]]
    es = {alan: next((b.index(a) for a in adlar if a in b), None) for alan, adlar in FORM_SUTUNLARI.items()}
    if es["kriter"] is None or es["puan"] is None:
        raise ValueError(f"{yol.name}: 'Kriter' ve 'Puan' sütunları gerekli")
    kriterler = []
    for i, r in enumerate(satirlar[1:], 1):
        al = lambda a: metin(r[es[a]]) if es[a] is not None and es[a] < len(r) else ""  # noqa: E731
        if not al("kriter"):
            continue
        try:
            puan = Decimal(al("puan").replace(",", "."))
        except InvalidOperation:
            raise ValueError(f"{yol.name}: '{al('kriter')}' kriterinin puanı okunamadı ('{al('puan')}')") from None
        kriterler.append(Kriter(al("kod") or f"K{i:02d}", al("bolum"), al("kriter"), al("aciklama"), puan, katla(al("kritik")) in ("evet", "e", "1", "x"),
                                [_normal(x) for x in al("ifadeler").split("|") if _normal(x)], katla(al("tur")).startswith("ihlal")))
    if not kriterler:
        raise ValueError(f"{yol.name}: kriter bulunamadı")
    return kriterler


@dataclass
class Gorusme:
    dosya: str
    meta: dict
    satirlar: list[tuple[str, str]]          # (konuşan: T / M / ?, maskeli metin)
    sonuclar: dict = field(default_factory=dict)
    ozet: dict = field(default_factory=dict)
    puan: Decimal | None = None
    ham_puan: Decimal | None = None
    kritik_hata: list[str] = field(default_factory=list)
    kontrol: list[str] = field(default_factory=list)

    @property
    def temsilci_metni(self) -> str:
        return "\n".join(m for k, m in self.satirlar if k == "T")

    @property
    def tam_metin(self) -> str:
        return "\n".join(f"{ {'T': 'Temsilci', 'M': 'Müşteri'}.get(k, 'Diğer')}: {m}" for k, m in self.satirlar)


def dokum_oku(yol: Path) -> Gorusme:
    ham = None
    for kod in ("utf-8-sig", "cp1254"):
        try:
            ham = yol.read_text(encoding=kod)
            break
        except UnicodeDecodeError:
            continue
    if ham is None:
        raise ValueError(f"{yol.name}: metin okunamadı")
    meta, satirlar = {}, []
    for satir in ham.splitlines():
        s = satir.strip()
        if not s:
            continue
        s = re.sub(r"^\[?\(?\d{1,2}:\d{2}(?::\d{2})?\)?\]?\s*", "", s)         # zaman damgası
        m = re.match(r"^([^:]{1,40}):\s*(.*)$", s)
        konusan = katla(m[1]) if m else ""
        if m and konusan in TEMSILCI_ADLARI:
            satirlar.append(("T", maskele(m[2])))
        elif m and konusan in MUSTERI_ADLARI:
            satirlar.append(("M", maskele(m[2])))
        elif m and not satirlar and konusan in ("temsilci adi", "tarih", "cagri no", "kuyruk", "konu", "sure"):
            meta[konusan] = m[2].strip()
        elif satirlar:
            satirlar[-1] = (satirlar[-1][0], satirlar[-1][1] + " " + maskele(s))
    if not any(k == "T" for k, _ in satirlar):
        raise ValueError(f"{yol.name}: 'Temsilci:' ile başlayan satır bulunamadı")
    return Gorusme(yol.name, meta, satirlar)


def konusma_paylari(g: Gorusme) -> dict:
    t = sum(len(m.split()) for k, m in g.satirlar if k == "T")
    mu = sum(len(m.split()) for k, m in g.satirlar if k == "M")
    return {"temsilci_kelime": t, "musteri_kelime": mu, "temsilci_pay": t / (t + mu) if t + mu else None,
            "tur": sum(1 for k, _ in g.satirlar if k == "T")}


# ----------------------------------------------------------------------------
# Model ve doğrulama
# ----------------------------------------------------------------------------

SEMA = {
    "type": "object",
    "properties": {
        "kriterler": {"type": "array", "items": {
            "type": "object",
            "properties": {"kod": {"type": "string"}, "sonuc": {"type": "string", "enum": SONUCLAR}, "gerekce": {"type": "string"}, "alinti": {"type": "string"}},
            "required": ["kod", "sonuc", "gerekce", "alinti"], "additionalProperties": False}},
        "guclu_yonler": {"type": "array", "items": {"type": "string"}},
        "gelisim_alanlari": {"type": "array", "items": {"type": "string"}},
        "geri_bildirim": {"type": "string"},
    },
    "required": ["kriterler", "guclu_yonler", "gelisim_alanlari", "geri_bildirim"], "additionalProperties": False,
}


def alinti_nerede(alinti: str, g: Gorusme) -> str:
    """'T' temsilci sözlerinde, 'M' yalnız müşteri sözlerinde, '' hiçbir yerde."""
    a = _normal(alinti)
    if len(a) < 6:
        return ""
    if any(a in _normal(m) for k, m in g.satirlar if k == "T"):
        return "T"
    return "M" if any(a in _normal(m) for k, m in g.satirlar if k != "T") else ""


def degerlendir(g: Gorusme, kriterler: list[Kriter], kritik_sifirla: bool = True) -> list[dict]:
    uy = []
    dokum = g.tam_metin
    if len(dokum) > AZAMI_DOKUM:
        dokum = dokum[:AZAMI_DOKUM]
        uy.append({"onem": "Orta", "tur": "Döküm kısaltıldı", "kim": g.dosya, "aciklama": f"İlk {AZAMI_DOKUM:,} karakter değerlendirildi"})
    mesaj = "\n".join(["<kalite_formu>", *[f'<kriter kod="{k.kod}" bolum="{k.bolum}" kritik="{"evet" if k.kritik else "hayir"}" tur="{"ihlal" if k.ihlal else "davranis"}">{k.ad}'
                                           + (f" — {k.aciklama}" if k.aciklama else "") + "</kriter>" for k in kriterler], "</kalite_formu>",
                       "<gorusme>", dokum, "</gorusme>"])
    y = llm.json_iste((BURASI / "prompt.md").read_text(encoding="utf-8"), mesaj, SEMA, max_tokens=8000)
    gelen = {}
    for s in y.get("kriterler", []):
        if s.get("kod") not in gelen:
            gelen[s.get("kod")] = s
    temsilci_n = _normal(g.temsilci_metni)
    for k in kriterler:
        s = gelen.get(k.kod)
        if not s or s.get("sonuc") not in SONUCLAR:
            g.sonuclar[k.kod] = {"sonuc": "yok", "gerekce": "", "alinti": "", "not": "Model bu kriter için sonuç döndürmedi"}
            uy.append({"onem": "Orta", "tur": "Kriter değerlendirilmedi", "kim": f"{g.dosya} · {k.kod}", "aciklama": f"{k.ad}: elle değerlendirin"})
            continue
        kayit = {"sonuc": s["sonuc"], "model_sonucu": s["sonuc"], "gerekce": s.get("gerekce", ""), "alinti": s.get("alinti", ""), "not": ""}
        yer = alinti_nerede(kayit["alinti"], g) if kayit["alinti"] else ""
        if k.ihlal:
            if kayit["sonuc"] in ("karsilanmadi", "kismen") and yer != "T":
                kayit["not"] = "İhlali gösteren alıntı temsilci sözlerinde bulunamadı"
                kayit["sonuc"], kayit["alinti"] = "inceleme", ""
            elif kayit["sonuc"] == "karsilandi" and kayit["alinti"] and not yer:
                kayit["alinti"] = ""
        elif kayit["sonuc"] in ("karsilandi", "kismen"):
            if yer != "T":
                kayit["not"] = "Alıntı müşterinin sözü; temsilci davranışını göstermiyor" if yer == "M" else "Alıntı dökümde bulunamadı"
                kayit["sonuc"], kayit["alinti"] = "inceleme", "" if not yer else kayit["alinti"]
        elif kayit["alinti"] and not yer:
            kayit["alinti"], kayit["not"] = "", "Alıntı dökümde bulunamadı; çıkarıldı"
        if k.ifadeler:
            var = any(i in temsilci_n for i in k.ifadeler)
            if kayit["sonuc"] == "karsilandi" and not var:
                kayit["sonuc"] = "inceleme"
                kayit["not"] = (kayit["not"] + "; " if kayit["not"] else "") + "Model karşılandı dedi ama anahtar ifadelerden hiçbiri temsilci sözlerinde yok"
            elif kayit["sonuc"] == "karsilanmadi" and var:
                kayit["not"] = (kayit["not"] + "; " if kayit["not"] else "") + "Anahtar ifade dökümde geçiyor; kontrol edin"
        if kayit["sonuc"] == "inceleme":
            uy.append({"onem": "Orta", "tur": "İnsan incelemesi gerekli", "kim": f"{g.dosya} · {k.kod}", "aciklama": f"{k.ad}: {kayit['not']}"})
        g.sonuclar[k.kod] = kayit
    for kod in gelen:
        if kod not in {k.kod for k in kriterler}:
            g.kontrol.append(f"Formda olmayan kriter kodu atıldı: {kod}")
    pay = sum((k.puan * KATSAYI[g.sonuclar[k.kod]["sonuc"]] for k in kriterler if g.sonuclar[k.kod]["sonuc"] in KATSAYI), Decimal(0))
    payda = sum((k.puan for k in kriterler if g.sonuclar[k.kod]["sonuc"] in KATSAYI), Decimal(0))
    g.ham_puan = (pay / payda * 100).quantize(Decimal("0.1"), ROUND_HALF_UP) if payda else None
    g.kritik_hata = [k.ad for k in kriterler if k.kritik and g.sonuclar[k.kod]["sonuc"] == "karsilanmadi"]
    g.puan = Decimal(0) if g.kritik_hata and kritik_sifirla and g.ham_puan is not None else g.ham_puan
    for ad in g.kritik_hata:
        uy.append({"onem": "Yüksek", "tur": "Kritik hata", "kim": g.dosya, "aciklama": f"Kritik kriter karşılanmadı: {ad}"
                   + ("; toplam puan 0" if kritik_sifirla else "")})
    for k in kriterler:
        if k.kritik and g.sonuclar[k.kod]["sonuc"] in ("inceleme", "yok"):
            uy.append({"onem": "Yüksek", "tur": "Kritik kriter belirsiz", "kim": f"{g.dosya} · {k.kod}", "aciklama": f"{k.ad}: sonuç doğrulanamadı; kaydı dinleyin"})
    g.ozet = {"guclu": [metin(x) for x in y.get("guclu_yonler", [])][:5], "gelisim": [metin(x) for x in y.get("gelisim_alanlari", [])][:5],
              "geri_bildirim": metin(y.get("geri_bildirim"))}
    return uy


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
SONUC_RENK = {"karsilandi": "E3F4E1", "kismen": "FFF4CE", "karsilanmadi": "FDE2E1", "inceleme": "EDE7FF", "yok": "EEEEEE"}
AI = PatternFill("solid", fgColor="EDE7FF")
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, basliklar, genislik=()):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def rapor_yaz(cikti: Path, s: dict) -> None:
    kriterler = s["kriterler"]
    wb = Workbook()
    oz = wb.active
    oz.title = "Özet"
    _baslik(oz, ["Döküm", "Temsilci", "Tarih", "Puan", "Ham Puan", "Kritik Hata", "İnceleme Bekleyen", "Temsilci Konuşma Payı", "Değerlendiren (insan onayı)"],
            (30, 20, 12, 8, 9, 40, 10, 12, 22))
    for g in s["gorusmeler"]:
        pay = konusma_paylari(g)
        bekleyen = sum(v["sonuc"] in ("inceleme", "yok") for v in g.sonuclar.values())
        oz.append([g.dosya, g.meta.get("temsilci adi", ""), g.meta.get("tarih", ""), None if g.puan is None else float(g.puan),
                   None if g.ham_puan is None else float(g.ham_puan), "; ".join(g.kritik_hata), bekleyen, pay["temsilci_pay"], ""])
        oz.cell(oz.max_row, 8).number_format = "0%"
        oz.cell(oz.max_row, 9).fill = PatternFill("solid", fgColor="FFF4CE")
        if g.kritik_hata:
            oz.cell(oz.max_row, 6).fill = PatternFill("solid", fgColor="FDE2E1")
    oz.append([])
    oz.append(["Puan = Σ ağırlık × (karşılandı 1, kısmen 0,5, karşılanmadı 0) / Σ ağırlık × 100; uygulanamaz ve inceleme bekleyen kriterler hariç."])
    oz.append([f"Model: {s['model']}. Sonuçlar taslaktır; kalite uzmanı kaydı dinleyerek onaylamalıdır."])

    ks = wb.create_sheet("Kriter Sonuçları")
    _baslik(ks, ["Döküm", "Kod", "Bölüm", "Kriter", "Ağırlık", "Kritik", "Sonuç", "Alınan Puan", "Gerekçe", "Alıntı (doğrulandı)", "Kontrol Notu", "Uzman Kararı"],
            (26, 6, 12, 40, 8, 7, 16, 8, 60, 50, 40, 16))
    for g in s["gorusmeler"]:
        for k in kriterler:
            v = g.sonuclar.get(k.kod, {"sonuc": "yok", "gerekce": "", "alinti": "", "not": ""})
            alinan = float(k.puan * KATSAYI[v["sonuc"]]) if v["sonuc"] in KATSAYI else None
            ks.append([g.dosya, k.kod, k.bolum, k.ad, float(k.puan), "Evet" if k.kritik else "", SONUC_ADI[v["sonuc"]], alinan, v["gerekce"], v["alinti"], v["not"], ""])
            renk = SONUC_RENK.get(v["sonuc"])
            if renk:
                ks.cell(ks.max_row, 7).fill = PatternFill("solid", fgColor=renk)
            for j in (9, 10):
                ks.cell(ks.max_row, j).alignment = UST
                ks.cell(ks.max_row, j).fill = AI
            ks.cell(ks.max_row, 11).alignment = UST
            ks.cell(ks.max_row, 12).fill = PatternFill("solid", fgColor="FFF4CE")
    ks.auto_filter.ref = f"A1:L{ks.max_row}"

    gb = wb.create_sheet("Geri Bildirim")
    _baslik(gb, ["Döküm", "Temsilci", "Puan", "Güçlü Yönler", "Gelişim Alanları", "Geri Bildirim Notu (taslak)"], (26, 20, 8, 50, 50, 90))
    for g in s["gorusmeler"]:
        gb.append([g.dosya, g.meta.get("temsilci adi", ""), None if g.puan is None else float(g.puan), "\n".join(f"• {x}" for x in g.ozet.get("guclu", [])),
                   "\n".join(f"• {x}" for x in g.ozet.get("gelisim", [])), g.ozet.get("geri_bildirim", "")])
        for j in (4, 5, 6):
            gb.cell(gb.max_row, j).alignment = UST
            gb.cell(gb.max_row, j).fill = AI

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Döküm / Kriter", "Açıklama"], (9, 26, 34, 100))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["kim"], u["aciklama"]])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(dokum_yolu: Path, form_yolu: Path, cikti: Path, *, kritik_sifirla: bool = True, evet: bool = False) -> dict:
    kriterler = form_oku(form_yolu)
    dosyalar = sorted(dokum_yolu.glob("*.txt")) if dokum_yolu.is_dir() else [dokum_yolu]
    if not dosyalar:
        raise llm.LLMHatasi(f"{dokum_yolu}: .txt döküm bulunamadı.")
    gorusmeler, uy = [], []
    for d in dosyalar:
        try:
            gorusmeler.append(dokum_oku(d))
        except ValueError as h:
            uy.append({"onem": "Yüksek", "tur": "Döküm okunamadı", "kim": d.name, "aciklama": str(h)})
    if not gorusmeler:
        raise llm.LLMHatasi("Değerlendirilecek döküm yok: " + "; ".join(u["aciklama"] for u in uy))
    print(f"[OK] {len(gorusmeler)} görüşme · {len(kriterler)} kriter ({sum(k.kritik for k in kriterler)} kritik) · toplam ağırlık {sum(k.puan for k in kriterler):g}")
    llm.onay_al(f"{len(gorusmeler)} görüşme dökümü (telefon, e-posta, TCKN, IBAN ve kart numaraları maskeli; adlar maskelenmez) ve kalite formu "
                "değerlendirme için gönderilecek.", evet)
    for g in gorusmeler:
        uy += degerlendir(g, kriterler, kritik_sifirla)
        pay = konusma_paylari(g)
        if pay["temsilci_pay"] is not None and pay["temsilci_pay"] > 0.75:
            uy.append({"onem": "Bilgi", "tur": "Konuşma payı", "kim": g.dosya, "aciklama": f"Temsilci konuşmanın %{pay['temsilci_pay'] * 100:.0f}'ini yapmış; dinleme dengesine bakın"})
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    uy.sort(key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))
    s = {"gorusmeler": gorusmeler, "kriterler": kriterler, "uyarilar": uy, "model": llm.kullanim_ozeti()}
    rapor_yaz(cikti, s)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Görüşme dökümünü kalite formundaki kriterlere göre puanlar, geri bildirim notu hazırlar.")
    p.add_argument("--dokum", type=Path, default=ORNEK / "gorusmeler", help="Döküm dosyası (.txt) veya klasörü; satırlar 'Temsilci:' / 'Müşteri:' ile başlar")
    p.add_argument("--form", type=Path, default=ORNEK / "kalite_formu.csv", help="Kod, Bölüm, Kriter, Açıklama, Puan, Kritik (Evet/Hayır), Anahtar İfadeler ('|' ile)")
    p.add_argument("--kritik-sifirlama-yok", action="store_true", help="Kritik hatada toplam puanı sıfırlama; yalnız işaretle")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "cagri_kalite_degerlendirmesi.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    for y in (a.dokum, a.form):
        if not y.exists():
            print(f"[X] Bulunamadı: {y}")
            return 1
    try:
        s = calistir(a.dokum, a.form, a.cikti, kritik_sifirla=not a.kritik_sifirlama_yok, evet=a.evet)
    except (llm.LLMHatasi, ValueError) as h:
        print(f"[X] {h}")
        return 1
    for g in s["gorusmeler"]:
        print(f"     {g.dosya:<30} puan {g.puan if g.puan is not None else '—':>5}" + (f"  KRİTİK HATA: {'; '.join(g.kritik_hata)}" if g.kritik_hata else ""))
    for u in s["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['kim']} · {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {s['model']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
