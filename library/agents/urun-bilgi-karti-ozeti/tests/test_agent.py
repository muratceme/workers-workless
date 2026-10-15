"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

TABLO = KLASOR / "ornek_veri" / "urunler.csv"
MESAJLAR = []


def od(fayda, ozellik, deger):
    return {"fayda": fayda, "ozellik": ozellik, "deger": deger}


def fk(diger, ozellik, bu, o, ifade):
    return {"diger_sku": diger, "ozellik": ozellik, "bu_deger": bu, "diger_deger": o, "ifade": ifade}


def kart(sku, dogru=True) -> dict:
    k = {"K-101": {"one_cikanlar": [od("Kutuyla 24 saate kadar dinleme", "Pil Süresi", "6 saat (kutuyla 24 saat)"), od("Kulakta hissedilmeyecek kadar hafif", "Ağırlık", "4,5 g"),
                                    od("Yağmur ve tere karşı dayanıklı", "Suya Dayanıklılık", "IPX4")],
                   "kisa_anlatim": "Günlük kullanım için hafif ve uygun fiyatlı bir kulaklık. Kutusuyla birlikte 24 saate kadar müzik dinleyebilirsiniz.",
                   "farklar": [fk("K-102", "Gürültü Engelleme", "Yok", "Aktif gürültü engelleme (ANC)", "Toplu taşımada gürültü engelleme isteyenler için Pro daha uygun.")],
                   "sik_sorular": [{"soru": "Suya dayanıklı mı?", "cevap": "Sıçrayan suya dayanıklıdır (IPX4); suya sokulmamalı.", "ozellik": "Suya Dayanıklılık"},
                                   {"soru": "Kablosuz şarj destekliyor mu?", "cevap": "Bilgi yok", "ozellik": ""}]},
         "K-102": {"one_cikanlar": [od("Gürültülü ortamda sessizlik", "Gürültü Engelleme", "Aktif gürültü engelleme (ANC)"),
                                    od("Kutuyla 30 saate kadar kullanım", "Pil Süresi", "8 saat (kutuyla 30 saat)"), od("Tere dayanıklı", "Suya Dayanıklılık", "IPX5")],
                   "kisa_anlatim": "Aktif gürültü engellemeli, uzun pil ömürlü bir kulaklık.",
                   "farklar": [fk("K-103", "Suya Dayanıklılık", "IPX5", "IP67", "Spor ve yoğun ter için IP67 sertifikalı Sport modeli daha uygun.")], "sik_sorular": []},
         "K-103": {"one_cikanlar": [od("Toza ve suya karşı dayanıklı", "Suya Dayanıklılık", "IP67"), od("Tek şarjla 9 saat", "Pil Süresi", "9 saat"),
                                    od("2 yıl garanti", "Garanti", "2 yıl")],
                   "kisa_anlatim": "Spor yapanlar için dayanıklı bir kulaklık.", "farklar": [], "sik_sorular": []},
         "S-201": {"one_cikanlar": [od("Tek şarjla 40 dakika temizlik", "Pil Süresi", "40 dakika"), od("Kolay taşınır", "Ağırlık", "2,6 kg"),
                                    od("Günlük temizliğe yeterli emiş", "Emiş Gücü", "150 AW")],
                   "kisa_anlatim": "Hafif gövdesiyle günlük temizlik için pratik; tek şarjla 40 dakika kullanım sunar.",
                   "farklar": [fk("S-202", "Emiş Gücü", "150 AW", "220 AW", "Halılı ve büyük evlerde daha yüksek emiş gücüyle V12 öne çıkar.")], "sik_sorular": []},
         "S-202": {"one_cikanlar": [od("Güçlü emiş", "Emiş Gücü", "220 AW"), od("60 dakika kullanım", "Pil Süresi", "60 dakika"), od("Geniş hazne", "Hazne Kapasitesi", "0,8 L")],
                   "kisa_anlatim": "Büyük evler için güçlü bir dikey süpürge.", "farklar": [], "sik_sorular": []}}[sku]
    k = dict(k, sku=sku)
    if not dogru:
        if sku == "K-101":
            k["one_cikanlar"] = [od("Uzun pil", "Pil Süresi", "8 saat")] + k["one_cikanlar"][1:]
            k["kisa_anlatim"] += " Tamamen su geçirmez."
        elif sku == "K-102":
            k["farklar"] = [dict(k["farklar"][0], diger_deger="IP68")]
        elif sku == "S-201":
            k["kisa_anlatim"] = k["kisa_anlatim"].replace("40 dakika", "45 dakika")
            k["farklar"] = k["farklar"] + [fk("K-101", "Ağırlık", "2,6 kg", "4,5 g", "Kulaklıktan ağır.")]
    return k


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    yazilacak = kullanici.split("</yazilacak_urunler>")[0]
    return {"kartlar": [kart(s, "<duzeltme>" in kullanici) for s in re.findall(r'<urun sku="([^"]+)"', yazilacak)]}


class KodTesti(unittest.TestCase):
    def test_okuma(self):
        u, sutun = agent.urunleri_oku(TABLO)
        self.assertEqual(len(u), 5)
        self.assertEqual(sutun[:3], ["Pil Süresi", "Gürültü Engelleme", "Suya Dayanıklılık"])
        self.assertNotIn("Emiş Gücü", u[0]["ozellikler"])
        self.assertTrue(agent._deger_ayni("6 saat", "6 saat (kutuyla 24 saat)"))
        self.assertFalse(agent._deger_ayni("2", "2,6 kg"))
        self.assertFalse(agent._deger_ayni("IP68", "IP67"))

    def test_denetim(self):
        u, _ = agent.urunleri_oku(TABLO)
        self.assertEqual(agent.denetle(u, {x["sku"]: kart(x["sku"]) for x in u}), [])
        s = agent.denetle(u, {x["sku"]: kart(x["sku"], False) for x in u})
        self.assertEqual(len(s), 5, s)
        self.assertTrue(any(x.startswith("K-101") and "'8 saat'" in x for x in s) and any("su gecirmez" in x for x in s)
                        and any(x.startswith("K-102") and "IP68" in x for x in s) and any(x.startswith("S-201") and "45" in x for x in s)
                        and any("başka kategoride" in x for x in s))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "u.xlsx"
            s = agent.calistir(TABLO, cikti, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            self.assertEqual(re.findall(r'<urun sku="([^"]+)"', MESAJLAR[1].split("</yazilacak_urunler>")[0]), ["K-101", "K-102", "S-201"])
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Kartlar", "Farklar"])
            self.assertEqual(wb["Farklar"].max_row, 4)
            self.assertIn("Örnek Kulaklık Air (K-101) — 1.899 TL", cikti.with_suffix(".txt").read_text(encoding="utf-8"))

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_kategori_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(TABLO, Path(tmp) / "u.xlsx", "Elektrikli Süpürge", tur=0, evet=True)
            self.assertEqual((len(s["urunler"]), len(s["kalan"])), (2, 2))
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--evet", "--kategori", "Telefon"]), 1)


if __name__ == "__main__":
    unittest.main()
