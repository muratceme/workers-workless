"""
Ürün Bilgi Kartı Özeti — Gerisi Bende AI Agent
Mağaza › Satış Danışmanı

1. Kod ürün tablosunu okur: SKU, Ürün, Kategori, Fiyat ve diğer her sütun bir teknik özellik (boş hücre = bilgi yok).
2. Model her ürün için müşteriye anlatılacak 3 öne çıkan fayda (özellik ve değeriyle), kısa anlatım, aynı kategorideki
   ürünlerle fark özeti ve sık sorulan sorulara cevap yazar.
3. Kod denetler:
   - Faydada, farkta ve cevapta adı geçen özellik tabloda var mı; yazılan değer tablodaki değerle aynı mı
   - Karşılaştırılan ürün aynı kategoride mi
   - Metindeki sayılar ürünün (farklarda karşılaştırılan ürünün de) özelliklerinde ya da fiyatında var mı
   - "Su geçirmez", "en iyi", "en ucuz", "ömür boyu" gibi iddialar ürünün özelliklerinde yazıyor mu
   Sorun varsa yalnız o ürünün kartı modele geri gönderilir (en fazla --tur kez).

Kullanım:
    python agent.py                                        # örnek kulaklık ve süpürgeler
    python agent.py --girdi urunler.xlsx --kategori "Kablosuz Kulaklık"
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
PAKET = 8
_OD = {"type": "object", "properties": {"fayda": S, "ozellik": S, "deger": S}, "required": ["fayda", "ozellik", "deger"], "additionalProperties": False}
SEMA = {
    "type": "object",
    "properties": {"kartlar": {"type": "array", "items": {"type": "object", "properties": {
        "sku": S, "one_cikanlar": {"type": "array", "items": _OD}, "kisa_anlatim": S,
        "farklar": {"type": "array", "items": {"type": "object", "properties": {"diger_sku": S, "ozellik": S, "bu_deger": S, "diger_deger": S, "ifade": S},
                                               "required": ["diger_sku", "ozellik", "bu_deger", "diger_deger", "ifade"], "additionalProperties": False}},
        "sik_sorular": {"type": "array", "items": {"type": "object", "properties": {"soru": S, "cevap": S, "ozellik": S},
                                                   "required": ["soru", "cevap", "ozellik"], "additionalProperties": False}}},
        "required": ["sku", "one_cikanlar", "kisa_anlatim", "farklar", "sik_sorular"], "additionalProperties": False}}},
    "required": ["kartlar"],
    "additionalProperties": False,
}
TEMEL = {"sku": ("SKU", "Stok Kodu", "Ürün Kodu"), "ad": ("Ürün", "Ürün Adı"), "kategori": ("Kategori",), "fiyat": ("Fiyat", "Satış Fiyatı", "Etiket Fiyatı")}
IDDIA = re.compile(r"\b(su gecirmez\w*|en iyi\w*|en ucuz\w*|en hafif\w*|en guclu\w*|en uzun\w*|rakipsiz|kirilmaz|omur boyu|sinirsiz|garantili olarak|"
                   r"tamamen sessiz|hic bozulmaz)\b")
BILGI_YOK = "Bilgi yok"


def deger_yaz(v) -> str:
    if v is None or v == "":
        return ""
    if isinstance(v, float):
        return sayi_yaz(v).replace(".", ",")
    return str(v).strip()


def urunleri_oku(yol: Path) -> tuple[list[dict], list[str]]:
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    bas = [str(x or "").strip() for x in s[0]]
    idx = {a: next((i for i, b in enumerate(bas) if katla(b) in {katla(e) for e in es}), None) for a, es in TEMEL.items()}
    if idx["sku"] is None or idx["ad"] is None:
        raise llm.LLMHatasi(f"{yol.name}: 'SKU' ve 'Ürün' sütunları gerekli. Başlıklar: {bas}")
    oz = [i for i, b in enumerate(bas) if b and i not in idx.values()]
    out = []
    for r in s[1:]:
        h = lambda i: deger_yaz(r[i]) if i is not None and i < len(r) else ""  # noqa: E731
        if not h(idx["sku"]):
            continue
        out.append({"sku": h(idx["sku"]), "ad": h(idx["ad"]), "kategori": h(idx["kategori"]), "fiyat": h(idx["fiyat"]),
                    "ozellikler": {bas[i]: h(i) for i in oz if h(i)}})
    return out, [bas[i] for i in oz]


def _ozellik(u: dict, ad: str) -> str | None:
    return next((v for k, v in u["ozellikler"].items() if katla(k) == katla(ad)), None)


def _deger_ayni(yazilan: str, gercek: str) -> bool:
    """Değer tablodakiyle aynı mı ya da tablodaki değerin bağımsız bir parçası mı ('6 saat' ↔ '6 saat (kutuyla 24 saat)')."""
    y, g = (" ".join(str(x).lower().split()) for x in (yazilan, gercek))
    if not y:
        return False
    i = g.find(y)
    while i >= 0:
        once, sonra = g[i - 1] if i else " ", g[i + len(y)] if i + len(y) < len(g) else " "
        if not re.match(r"[\w.,]", once) and not re.match(r"[\w.,]", sonra):
            return True
        i = g.find(y, i + 1)
    return False


def kaynak_metin(u: dict) -> str:
    return " ".join([u["ad"], u["fiyat"]] + [f"{k} {v}" for k, v in u["ozellikler"].items()])


def denetle(urunler: list[dict], sonuc: dict) -> list[str]:
    sorun, sku = [], {u["sku"]: u for u in urunler}
    for u in urunler:
        a, k = u["sku"], sonuc.get(u["sku"])
        if not k:
            sorun.append(f"{a}: kart yazılmadı")
            continue
        kaynak = kaynak_metin(u)

        def ozellik_kontrol(yer, ad, deger, urun):
            g = _ozellik(urun, ad)
            if g is None:
                if katla(ad) == "fiyat" and _deger_ayni(deger, urun["fiyat"]):
                    return
                sorun.append(f"{a}: {yer}: '{ad}' özelliği {urun['sku']} için tabloda yok")
            elif deger and not _deger_ayni(deger, g):
                sorun.append(f"{a}: {yer}: {urun['sku']} '{ad}' değeri tabloda '{g}', yazılan '{deger}'")

        if len(k["one_cikanlar"]) < 3:
            sorun.append(f"{a}: 3 öne çıkan fayda istendi, {len(k['one_cikanlar'])} yazıldı")
        metinler = [("Anlatım", k["kisa_anlatim"], kaynak)]
        for i, o in enumerate(k["one_cikanlar"], 1):
            ozellik_kontrol(f"fayda {i}", o["ozellik"], o["deger"], u)
            metinler.append((f"Fayda {i}", o["fayda"], kaynak))
        for i, f in enumerate(k["farklar"], 1):
            d = sku.get(f["diger_sku"])
            if not d:
                sorun.append(f"{a}: fark {i}: {f['diger_sku']} tabloda yok")
                continue
            if katla(d["kategori"]) != katla(u["kategori"]):
                sorun.append(f"{a}: fark {i}: {d['sku']} başka kategoride ({d['kategori']}); aynı kategorideki ürünlerle karşılaştırın")
            ozellik_kontrol(f"fark {i}", f["ozellik"], f["bu_deger"], u)
            ozellik_kontrol(f"fark {i}", f["ozellik"], f["diger_deger"], d)
            metinler.append((f"Fark {i}", f["ifade"], kaynak + " " + kaynak_metin(d)))
        for i, q in enumerate(k["sik_sorular"], 1):
            if q["ozellik"] and _ozellik(u, q["ozellik"]) is None and katla(q["ozellik"]) != "fiyat":
                sorun.append(f"{a}: soru {i}: '{q['ozellik']}' özelliği tabloda yok; bilgi yoksa ozellik boş, cevap '{BILGI_YOK}' olmalı")
            metinler.append((f"Soru {i}", q["cevap"], kaynak))
        for yer, m, src in metinler:
            fazla = sorted(sayilar(m) - sayilar(src))
            if fazla:
                sorun.append(f"{a}: {yer}: tabloda olmayan sayı: {', '.join(fazla)}")
            for x in IDDIA.finditer(katla(m)):
                if x.group(0) not in katla(src):
                    sorun.append(f"{a}: {yer}: '{x.group(0)}' iddiası ürün özelliklerinde yok")
    return sorun


def urun_satiri(u: dict) -> str:
    oz = "; ".join(f"{k}: {v}" for k, v in u["ozellikler"].items())
    return f"<urun sku=\"{u['sku']}\" kategori=\"{u['kategori']}\">{u['ad']} | Fiyat: {u['fiyat'] or '-'} | {oz}</urun>"


def uret(urunler: list[dict], duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    sonuc, hedef = {}, urunler
    if duzeltme is not None:  # düzeltmede yalnız sorunlu kartlar yeniden yazılır
        sorunlu = {s.split(":")[0] for s in duzeltme}
        sonuc.update({a: b for a, b in (onceki or {}).items() if a not in sorunlu})
        hedef = [u for u in urunler if u["sku"] in sorunlu]
    for j in range(0, len(hedef), PAKET):
        parca = hedef[j:j + PAKET]
        ids = {u["sku"] for u in parca}
        kategoriler = {u["kategori"] for u in parca}
        karsilastirma = [u for u in urunler if u["kategori"] in kategoriler and u["sku"] not in ids]
        d = [s for s in (duzeltme or []) if s.split(":")[0] in ids]
        mesaj = ("<yazilacak_urunler>\n" + "\n".join(urun_satiri(u) for u in parca) + "\n</yazilacak_urunler>\n<ayni_kategorideki_diger_urunler>\n"
                 + ("\n".join(urun_satiri(u) for u in karsilastirma) or "(yok)") + "\n</ayni_kategorideki_diger_urunler>"
                 + duzeltme_blogu(d, [v for a, v in (onceki or {}).items() if a in ids]))
        for k in llm.json_iste(sistem, mesaj, SEMA).get("kartlar", []):
            if k.get("sku") in ids:
                sonuc[k["sku"]] = k
    return sonuc


def rapor_yaz(cikti: Path, urunler, sonuc, kalan, tur) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kartlar"
    baslik(ws, ["SKU", "Ürün", "Kategori", "Fiyat", "Öne Çıkanlar", "Kısa Anlatım", "Sık Sorulanlar", "Kontrol"], (10, 28, 18, 11, 55, 60, 60, 30))
    for u in urunler:
        k = sonuc.get(u["sku"])
        sor = [s.split(": ", 1)[1] for s in kalan if s.startswith(u["sku"] + ":")]
        if not k:
            ws.append([u["sku"], u["ad"], u["kategori"], u["fiyat"], "", "AI yanıt vermedi"])
            continue
        ws.append([u["sku"], u["ad"], u["kategori"], u["fiyat"], "\n".join(f"• {o['fayda']} ({o['ozellik']}: {o['deger']})" for o in k["one_cikanlar"]),
                   k["kisa_anlatim"], "\n".join(f"S: {q['soru']}\nC: {q['cevap']}" for q in k["sik_sorular"]), "; ".join(sor) or "Geçti"])
        boya(ws, [5, 6, 7], AI_DOLGU)
        boya(ws, [8], KIRMIZI if sor else YESIL)
    son_not(ws, "Önemli: Kartları mağazada kullanmadan önce üretici teknik belgeleri ve güncel etiket fiyatıyla karşılaştırın.", kalan, tur)
    hizala(ws)
    f = wb.create_sheet("Farklar")
    baslik(f, ["SKU", "Karşılaştırılan", "Özellik", "Bu Ürün", "Diğer Ürün", "Müşteriye Anlatım"], (10, 14, 20, 22, 22, 70))
    for u in urunler:
        for x in (sonuc.get(u["sku"]) or {}).get("farklar", []):
            f.append([u["sku"], x["diger_sku"], x["ozellik"], x["bu_deger"], x["diger_deger"], x["ifade"]])
            boya(f, [6], AI_DOLGU)
    hizala(f)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    kart = []
    for u in urunler:
        k = sonuc.get(u["sku"])
        if not k or any(s.startswith(u["sku"] + ":") for s in kalan):
            continue
        kart += [f"{u['ad']} ({u['sku']}) — {u['fiyat']}", k["kisa_anlatim"]] + [f"  • {o['fayda']}" for o in k["one_cikanlar"]]
        kart += [f"  ↔ {x['diger_sku']}: {x['ifade']}" for x in k["farklar"]] + [""]
    cikti.with_suffix(".txt").write_text("\n".join(kart), encoding="utf-8")


def calistir(girdi: Path, cikti: Path, kategori: str = "", tur: int = 2, evet: bool = False) -> dict:
    urunler, sutunlar = urunleri_oku(girdi)
    if kategori:
        urunler = [u for u in urunler if katla(u["kategori"]) == katla(kategori)]
    if not urunler:
        raise llm.LLMHatasi(f"{girdi.name}: ürün yok{' (kategori: ' + kategori + ')' if kategori else ''}.")
    print(f"[OK] {len(urunler)} ürün · {len({u['kategori'] for u in urunler})} kategori · {len(sutunlar)} özellik sütunu")
    llm.onay_al(f"{len(urunler)} ürünün adı, fiyatı ve teknik özellikleri gönderilecek.", evet)
    sonuc, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(urunler, d, o), lambda x: denetle(urunler, x), tur)
    rapor_yaz(cikti, urunler, sonuc, kalan, yapilan)
    return {"urunler": urunler, "sonuc": sonuc, "kalan": kalan, "tur": yapilan}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    p = argparse.ArgumentParser(description="Ürün teknik bilgilerinden müşteriye anlatılacak kısa özellik ve fark özetleri hazırlar.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "urunler.csv",
                   help="Ürün tablosu (.xlsx/.csv): SKU, Ürün; Kategori, Fiyat; diğer her sütun bir teknik özellik")
    p.add_argument("--kategori", default="", help="Yalnız bu kategorideki ürünler")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "urun_kartlari.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.cikti, a.kategori, a.tur, a.evet)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(s['sonuc'])}/{len(s['urunler'])} ürün kartı")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()} (+ {a.cikti.with_suffix('.txt').name})")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
