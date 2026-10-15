"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

RAPOR = KLASOR / "ornek_veri" / "trend_notlari.txt"
BRIF = KLASOR / "ornek_veri" / "koleksiyon_brifi.txt"
MESAJLAR = []


def m(kategori, trend, alinti, satir, uygulama, gruplar, oncelik="Yüksek"):
    return {"kategori": kategori, "trend": trend, "alinti": alinti, "satirlar": [satir], "uygulama": uygulama, "urun_gruplari": gruplar, "oncelik": oncelik}


def ozet(dogru=True) -> dict:
    t = {"ozet": "Toprak tonları, keten ve el işi detaylarla rahat ve doğal bir yaz koleksiyonu.",
         "maddeler": [m("Renk", "Güneşte solmuş toprak tonları", "terrakota, kum beji ve zeytin yeşili ana renkler", 2,
                        "Gömlek, pantolon ve elbiselerde ana renkler olarak kullanılır.", ["Gömlek", "Pantolon", "Elbise"]),
                      m("Renk", "Mercan vurgu rengi", "canlı bir mercan tonu (örnek kod: 16-1546 TCX)", 3, "Tişört ve elbiselerde kapsül renk; 16-1546 TCX referans.",
                        ["Tişört", "Elbise"], "Orta"),
                      m("Kumaş", "Keten ve keten-pamuk", "keten ve keten-pamuk karışımları yaygınlaşıyor", 4, "Gömlek ve pantolonlarda yıkanmış keten.",
                        ["Gömlek", "Pantolon", "Elbise"]),
                      m("Kumaş", "Sertifikalı organik pamuk", "sertifikalı organik pamuk", 5, "Tişörtlerde organik pamuk.", ["Tişört"], "Orta"),
                      m("Silüet", "Bol kesim ve geniş paça", "düşük omuzlu gömlekler ve geniş paça pantolonlar", 6, "Gömleklerde düşük omuz, pantolonlarda geniş paça.",
                        ["Gömlek", "Pantolon"]),
                      m("Detay", "El işi nakış ve doğal düğme", "kontrast dikişler", 7, "Gömlek ve elbiselerde kontrast dikiş ve ahşap düğme.", ["Gömlek", "Elbise"], "Orta"),
                      m("Desen / Baskı", "Küçük çiçek baskı", "küçük ölçekli çiçek baskılar", 8, "Elbise ve eteklerde küçük çiçek baskı.", ["Elbise", "Etek"])],
         "uygulanmayacaklar": [{"trend": "Hasır ve rafya çantalar", "neden": "Aksesuar koleksiyon ürün gruplarında yok.", "satirlar": [9]},
                               {"trend": "Metalik ve payetli abiye kumaşlar", "neden": "Brifte kaçınılacaklar arasında.", "satirlar": [10]}]}
    if not dogru:
        t["maddeler"][0]["alinti"] = "terrakota ve kum beji ana renkler"
        t["maddeler"][1]["uygulama"] = "Tişörtlerde 17-1463 TCX mercan."
        t["maddeler"][2]["urun_gruplari"] = ["Gömlek", "Ceket"]
        del t["maddeler"][4]
        t["maddeler"].append(m("Kumaş", "Metalik kumaş ve payet", "metalik ve payetli kumaşlar", 10, "Elbiselerde payet detay.", ["Elbise"], "Düşük"))
    return t


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return ozet("<duzeltme>" in kullanici)


class KodTesti(unittest.TestCase):
    def test_okuma(self):
        b, l = agent.brifi_oku(BRIF), agent.raporu_oku(RAPOR)
        self.assertEqual(b["grup_listesi"], ["Gömlek", "Pantolon", "Elbise", "Etek", "Tişört"])
        self.assertEqual(b["kacin_listesi"], ["Payet", "metalik kumaş", "deri"])
        self.assertEqual(agent.raporda_kategoriler(l), ["Renk", "Kumaş", "Silüet", "Detay", "Desen / Baskı"])
        self.assertEqual(agent.raporda_kategoriler(["Sezonun rengi zeytin yeşili."]), ["Renk"])

    def test_denetim(self):
        b, l = agent.brifi_oku(BRIF), agent.raporu_oku(RAPOR)
        self.assertEqual(agent.denetle(ozet(), l, b), [])
        s = agent.denetle(ozet(False), l, b)
        self.assertEqual(len(s), 6, s)
        self.assertTrue(any("alıntı" in x for x in s) and any("17-1463 TCX" in x for x in s) and any("Ceket" in x for x in s)
                        and any("'Silüet'" in x for x in s) and any("'Payet'" in x for x in s) and any("'metalik kumaş'" in x for x in s))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            s = agent.calistir(RAPOR, BRIF, cikti, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Maddeler", "Uygulanmayacaklar"])
            self.assertEqual([r[0] for r in wb["Maddeler"].iter_rows(min_row=2, max_row=8, values_only=True)],
                             ["Renk", "Renk", "Kumaş", "Kumaş", "Silüet", "Detay", "Desen / Baskı"])
            self.assertIn("UYGULANMAYACAKLAR", cikti.with_suffix(".txt").read_text(encoding="utf-8"))

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(RAPOR, BRIF, Path(tmp) / "t.xlsx", tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 6)
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--evet", "--brif", str(BRIF.with_name("yok.txt"))]), 1)


if __name__ == "__main__":
    unittest.main()
