"""
Trend Raporu Özeti — Workers / Workless AI Agent
Tasarım › Moda Tasarımcısı

1. Kod sezon trend raporunu satır (paragraf) satır numaralar ve koleksiyon brifini okur (sezon, hedef müşteri, ürün
   grupları, fiyat segmenti, ana kumaşlar, kaçınılacaklar).
2. Model raporu koleksiyona uygulanabilir renk, kumaş, silüet, detay ve desen / baskı maddelerine indirir: her madde rapordan
   aynen alıntı, uygulanacağı ürün grupları ve öncelikle gelir; koleksiyona uymayan trendler gerekçesiyle ayrılır.
3. Kod denetler:
   - Alıntı dayandığı rapor satırında aynen geçiyor mu; satır numaraları geçerli mi
   - Ürün grupları brifte var mı; Pantone (TCX / TPX) kodu ve sayılar raporda / brifte var mı
   - Raporda geçen bir kategori (renk, kumaş, silüet, detay, desen) atlanmış mı
   - Brifteki "kaçınılacaklar" (ör. payet, deri) bir uygulama maddesinde önerilmiş mi
   Sorun varsa modele geri gönderilir (en fazla --tur kez).

Kullanım:
    python agent.py                                              # örnek (kurgusal) trend notları + brif
    python agent.py --rapor trend.pdf --brif koleksiyon_brifi.txt
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
K = {"type": "array", "items": {"type": "integer"}}
KATEGORILER = ["Renk", "Kumaş", "Silüet", "Detay", "Desen / Baskı"]
SEMA = {
    "type": "object",
    "properties": {
        "ozet": S,
        "maddeler": {"type": "array", "items": {"type": "object", "properties": {
            "kategori": {"type": "string", "enum": KATEGORILER}, "trend": S, "alinti": S, "satirlar": K, "uygulama": S,
            "urun_gruplari": {"type": "array", "items": S}, "oncelik": {"type": "string", "enum": ["Yüksek", "Orta", "Düşük"]}},
            "required": ["kategori", "trend", "alinti", "satirlar", "uygulama", "urun_gruplari", "oncelik"], "additionalProperties": False}},
        "uygulanmayacaklar": {"type": "array", "items": {"type": "object", "properties": {"trend": S, "neden": S, "satirlar": K},
                                                         "required": ["trend", "neden", "satirlar"], "additionalProperties": False}},
    },
    "required": ["ozet", "maddeler", "uygulanmayacaklar"],
    "additionalProperties": False,
}
ANAHTAR = {"Renk": r"\b(renk\w*|reng\w*|ton\w*|palet\w*|tcx|tpx)\b", "Kumaş": r"\b(kumas\w*|keten\w*|pamuk\w*|polyester\w*|viskon\w*|iplik\w*|orme\w*|dokuma\w*|denim\w*)\b",
           "Silüet": r"\b(siluet\w*|kesim\w*|oversize|paca\w*|omuz\w*)\b", "Detay": r"\b(detay\w*|nakis\w*|dikis\w*|dugme\w*|fermuar\w*)\b",
           "Desen / Baskı": r"\b(desen\w*|baski\w*|motif\w*|cicek\w*|cizgi\w*|cizgili|kareli)\b"}
PANTONE = re.compile(r"\b\d{2}-\d{4}\s*(?:TCX|TPX|TPG)\b", re.I)
BRIF_ALANLARI = {"sezon": "Sezon", "hedef": "Hedef Müşteri", "gruplar": "Ürün Grupları", "segment": "Fiyat Segmenti", "kumaslar": "Ana Kumaşlar",
                 "kacinilacak": "Kaçınılacaklar"}


def brifi_oku(yol: Path) -> dict:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: koleksiyon brifi bulunamadı.")
    b = {a: "" for a in BRIF_ALANLARI}
    for s in belge.metin_oku(yol).splitlines():
        a, _, d = s.partition(":")
        k = next((x for x, v in BRIF_ALANLARI.items() if katla(v) == katla(a)), None)
        if k:
            b[k] = d.strip()
    b["grup_listesi"] = [x.strip() for x in re.split(r"[,;]", b["gruplar"]) if x.strip()]
    b["kacin_listesi"] = [x.strip() for x in re.split(r"[,;]", b["kacinilacak"]) if x.strip()]
    if not b["grup_listesi"]:
        raise llm.LLMHatasi(f"{yol.name}: 'Ürün Grupları: Gömlek, Pantolon, …' satırı gerekli.")
    return b


def raporu_oku(yol: Path) -> list[str]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: trend raporu bulunamadı.")
    return [s.strip() for s in belge.metin_oku(yol).splitlines() if s.strip()]


def raporda_kategoriler(l: list[str]) -> list[str]:
    k = katla(" ".join(l))
    return [c for c, d in ANAHTAR.items() if re.search(d, k)]


def denetle(t: dict, l: list[str], b: dict) -> list[str]:
    n, sorun = len(l), []
    rapor = "\n".join(l)
    izinli = sayilar(rapor) | sayilar(" ".join(b[a] for a in BRIF_ALANLARI))
    gruplar = {katla(g) for g in b["grup_listesi"]}
    for i, m in enumerate(t["maddeler"], 1):
        ad = f"Madde {i} ({m['kategori']})"
        ks = m["satirlar"]
        if not ks or [k for k in ks if not 1 <= k <= n]:
            sorun.append(f"{ad}: dayanak satır yok ya da geçersiz {[k for k in ks if not 1 <= k <= n]}")
        kaynak = " ".join(l[k - 1] for k in ks if 1 <= k <= n)
        if not metinde_var(m["alinti"], kaynak):
            sorun.append(f"{ad}: alıntı dayandığı satırda aynen yok: '{m['alinti'][:60]}'")
        yanlis = [g for g in m["urun_gruplari"] if katla(g) not in gruplar]
        if yanlis or not m["urun_gruplari"]:
            sorun.append(f"{ad}: brifte olmayan ürün grubu: {', '.join(yanlis) or '(boş)'}")
        metin = " ".join((m["trend"], m["alinti"], m["uygulama"]))
        for kod in PANTONE.findall(metin):
            if katla(kod) not in katla(rapor):
                sorun.append(f"{ad}: Pantone kodu {kod} raporda yok (renk kodu uydurulmamalı)")
        fazla = sorted(sayilar(PANTONE.sub(" ", metin)) - izinli)
        if fazla:
            sorun.append(f"{ad}: raporda ve brifte olmayan sayı: {', '.join(fazla)}")
        for x in b["kacin_listesi"]:
            if katla(x) and katla(x) in katla(m["trend"] + " " + m["uygulama"]):
                sorun.append(f"{ad}: brifte kaçınılacaklar arasında olan '{x}' önerilmiş (uygulanmayacaklara taşıyın)")
    for i, u in enumerate(t["uygulanmayacaklar"], 1):
        if not u["satirlar"] or [k for k in u["satirlar"] if not 1 <= k <= n]:
            sorun.append(f"Uygulanmayacak {i}: dayanak satır yok ya da geçersiz")
    kapsanan = {m["kategori"] for m in t["maddeler"]}
    for c in raporda_kategoriler(l):
        if c not in kapsanan:
            sorun.append(f"Genel: raporda '{c}' trendleri var ama bu kategoride madde yok")
    return sorun


def uret(l: list[str], b: dict, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    rapor = "\n".join(f"{i}. {s}" for i, s in enumerate(l, 1))
    brif = "\n".join(f"{v}: {b[a]}" for a, v in BRIF_ALANLARI.items() if b[a])
    mesaj = f"<koleksiyon_brifi>\n{brif}\n</koleksiyon_brifi>\n<trend_raporu>\n{rapor}\n</trend_raporu>" + duzeltme_blogu(duzeltme, onceki)
    return llm.json_iste(sistem, mesaj, SEMA)


def rapor_yaz(cikti: Path, t, l, b, kalan, tur) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Maddeler"
    baslik(ws, ["Kategori", "Trend", "Rapordan Alıntı", "Satır", "Koleksiyona Uygulama", "Ürün Grupları", "Öncelik", "Karar"],
           (13, 34, 50, 7, 60, 24, 9, 10))
    renk = {"Yüksek": YESIL, "Orta": SARI, "Düşük": None}
    for m in sorted(t["maddeler"], key=lambda m: KATEGORILER.index(m["kategori"])):
        ws.append([m["kategori"], m["trend"], m["alinti"], ", ".join(map(str, m["satirlar"])), m["uygulama"], ", ".join(m["urun_gruplari"]), m["oncelik"], ""])
        boya(ws, [2, 5], AI_DOLGU)
        if renk[m["oncelik"]]:
            boya(ws, [7], renk[m["oncelik"]])
        boya(ws, [8], SARI)
    son_not(ws, "Önemli: Maddeler tasarım ekibinin değerlendirmesi içindir. Renk kodlarını fiziksel renk kartlarıyla, kumaşları tedarikçi "
                "numuneleriyle doğrulayın.", kalan, tur)
    hizala(ws)
    u = wb.create_sheet("Uygulanmayacaklar")
    baslik(u, ["Trend", "Neden", "Satır"], (40, 70, 8))
    for x in t["uygulanmayacaklar"]:
        u.append([x["trend"], x["neden"], ", ".join(map(str, x["satirlar"]))])
        boya(u, [2], AI_DOLGU)
    hizala(u)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    metin = [f"TREND ÖZETİ — {b['sezon']}", "", t["ozet"], ""]
    for c in KATEGORILER:
        ms = [m for m in t["maddeler"] if m["kategori"] == c]
        if ms:
            metin += [c.upper()] + [f"• {m['trend']} → {m['uygulama']} ({', '.join(m['urun_gruplari'])}; {m['oncelik']})" for m in ms] + [""]
    if t["uygulanmayacaklar"]:
        metin += ["UYGULANMAYACAKLAR"] + [f"• {x['trend']}: {x['neden']}" for x in t["uygulanmayacaklar"]]
    cikti.with_suffix(".txt").write_text("\n".join(metin), encoding="utf-8")


def calistir(rapor_yolu: Path, brif_yolu: Path, cikti: Path, tur: int = 2, evet: bool = False) -> dict:
    l = raporu_oku(rapor_yolu)
    b = brifi_oku(brif_yolu)
    print(f"[OK] {len(l)} rapor satırı · kategoriler: {', '.join(raporda_kategoriler(l))} · ürün grupları: {', '.join(b['grup_listesi'])}")
    llm.onay_al(f"{len(l)} satırlık trend raporu ve koleksiyon brifi gönderilecek. Lisanslı trend raporlarının üçüncü kişilerle "
                "paylaşımına ilişkin sözleşme koşullarınızı kontrol edin.", evet)
    t, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(l, b, d, o), lambda x: denetle(x, l, b), tur)
    rapor_yaz(cikti, t, l, b, kalan, yapilan)
    return {"satirlar": l, "brif": b, "taslak": t, "kalan": kalan, "tur": yapilan}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    o = BURASI / "ornek_veri"
    p = argparse.ArgumentParser(description="Sezon trend raporlarını koleksiyona uygulanabilir renk, kumaş ve detay maddelerine indirir.")
    p.add_argument("--rapor", type=Path, default=o / "trend_notlari.txt", help="Trend raporu (.pdf/.docx/.txt)")
    p.add_argument("--brif", type=Path, default=o / "koleksiyon_brifi.txt",
                   help="Koleksiyon brifi (.txt/.docx): Sezon, Hedef Müşteri, Ürün Grupları, Fiyat Segmenti, Ana Kumaşlar, Kaçınılacaklar")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "trend_ozeti.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.rapor, a.brif, a.cikti, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(s['taslak']['maddeler'])} uygulama maddesi · {len(s['taslak']['uygulanmayacaklar'])} uygulanmayacak trend")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()} (+ {a.cikti.with_suffix('.txt').name})")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
