"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

DOSYA = KLASOR / "ornek_veri" / "ornek_dahili_su"
MESAJLAR = []
P, E = "police_ozel_sartlar.txt", "eksper_on_raporu.txt"


def a(k, t):
    return {"kaynak": k, "alinti": t}


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return {
        "olay_ozeti": "[SİGORTALI] deposunda tavandaki su borusu patladı; kırtasiye emtiası ıslandı.",
        "teminat_degerlendirmesi": [{"teminat": "Dahili su", "sonuc": "Kapsamda", "gerekce": "Boru patlaması dahili su teminatına girer.",
                                     "police_alintilari": [a(P, "su, kalorifer ve kanalizasyon tesisatının patlaması, taşması, sızması veya kırılması")],
                                     "olay_alintilari": [a(E, "su tesisatı borusunun patlaması sonucu depoya su dolmuştur")]}],
        "istisnalar": [
            {"istisna": "Bakımsızlık / paslanma (2.3)", "uygulanir": "Belirsiz", "gerekce": "Korozyon var ama ani patlama.",
             "police_alintilari": [a(P, "Tesisatın bakımsızlığı, paslanma, çürüme veya aşınma sonucu zamanla oluşan sızıntılardan doğan zararlar teminat dışıdır.")],
             "olay_alintilari": [a(E, "Patlayan borunun dış yüzeyinde yaygın korozyon görülmüştür.")]},
            {"istisna": "Zeminde istif (3.1)", "uygulanir": "Uygulanır", "gerekce": "Koliler zemindeydi.",
             "police_alintilari": [a(P, "zemin ile temas eden emtiada meydana gelen su zararları teminat dışıdır")],
             "olay_alintilari": [a(E, "bir kısmı zemin üzerinde duruyordu")]},
        ],
        "kapsam_disi_kalemler": [
            {"aciklama": "Zemin üzerindeki koliler", "tutar": 70000, "olay_alintisi": "bunun 70.000 TL'si zemin üzerindeki kolilere aittir"},
            {"aciklama": "Kurutma gideri", "tutar": 15000, "olay_alintisi": "Kurutma gideri 15.000 TL."},
        ],
        "sonuc": {"degerlendirme": "Kapsamda", "gerekce": "Tazminat yaklaşık 95.000 TL olabilir."},
        "ek_bilgi_gerekenler": ["Tesisat bakım kayıtları."],
        "celiskiler": ["Sigortalı zararı 150.000 TL beyan etmiş, eksper 120.000 TL tespit etmiş."],
    }


class KontrolTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.b, cls.t, cls.m, cls.a = agent.dosyalari_oku(DOSYA)
        cls.k = agent.kontrol_et(cls.b, cls.t, cls.m)

    def test_okuma(self):
        self.assertEqual({d.ad: d.tur for d in self.m}, {E: "Olay belgesi", "hasar_bilgileri.txt": "Hasar bilgileri", "ihbar.txt": "Olay belgesi",
                                                        P: "Poliçe / şartlar"})
        self.assertEqual(self.k["teminat"]["teminat"], "Dahili su")
        self.assertEqual(self.k["oran"], D(300000) / D(400000))

    def test_muafiyet(self):
        m = agent.muafiyet_coz("%5 en az 1.000 en fazla 20.000")
        self.assertEqual((m["oran"], m["en_az"], m["en_cok"]), (D("0.05"), D(1000), D(20000)))
        self.assertEqual(agent.muafiyet_tutari(m, D(1000000)), D(20000))
        self.assertEqual(agent.muafiyet_tutari(agent.muafiyet_coz("10.000 TL"), D(5)), D(10000))

    def test_on_hesap(self):
        self.assertEqual([(x, int(v)) for x, v in self.k["hesap"]["adimlar"]],
                         [("Hasar tutarı (eksper)", 120000), ("Eksik sigorta oranı (%75,0)", -30000), ("Muafiyet (%2, en az 5.000 TL)", -5000),
                          ("Ön hesap tazminat", 85000)])
        tem = {"teminat": "X", "limit": D(50000), "muafiyet": agent.muafiyet_coz("")}
        self.assertEqual(agent.on_hesap(D(80000), None, tem)["sonuc"], D(50000))

    def test_bulgular(self):
        metin = " ".join(f"{o} {a}" for o, a, _ in self.k["bulgular"])
        self.assertIn("orta Eksik sigorta: sigorta bedeli 300.000 TL, gerçek değer 400.000 TL (oran %75,0)", metin)
        self.assertIn("İhbar hasardan 8 gün sonra", metin)
        self.assertNotIn("yüksek", metin)

    def test_sure_disi_ve_prim(self):
        b = {"metin": "Poliçe Başlangıç: 01.06.2026\nPoliçe Bitiş: 01.06.2027\nHasar Tarihi: 10.06.2026\nPrim Durumu: Ödenmedi\n"
                      "Talep Edilen Teminat: Deprem\n", "dosya": "x"}
        metin = " ".join(a for _, a, _ in agent.kontrol_et(b, self.t, self.m)["bulgular"])
        self.assertIn("başlangıcından 9 gün sonra", metin)
        self.assertIn("Prim durumu: Ödenmedi", metin)
        self.assertIn("'Deprem' poliçe teminat tablosunda yok", metin)
        b["metin"] = b["metin"].replace("10.06.2026", "05.06.2027")
        self.assertIn("poliçe süresi dışında", " ".join(a for _, a, _ in agent.kontrol_et(b, self.t, self.m)["bulgular"]))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_uctan_uca(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            s = agent.calistir(DOSYA, cikti, evet=True)
            m = MESAJLAR[0]
            for gizli in ("Örnek Kırtasiye", "IP-2026-77881", "Burcu Örnekçi", "0532 000 11 22"):
                self.assertNotIn(gizli, m)
            for var in ("[SİGORTALI]", "[POLİÇE NO]", "[KİŞİ-1]", "<police_metinleri>", "Ön hesap adımı: Ön hesap tazminat 85.000 TL"):
                self.assertIn(var, m)
            y = s["yanit"]
            self.assertEqual([x["dogrulandi"] for i in y["istisnalar"] for x in i["olay_alintilari"]], [True, False])
            self.assertEqual([x["dogrulandi"] for x in y["kapsam_disi_kalemler"]], [True, False])
            self.assertEqual(y["ikinci_hesap"]["sonuc"], D(32500))                  # (120.000 − 70.000) × 0,75 − 5.000
            sorun = " ".join(a for _, a, _ in y["sorunlar"])
            self.assertIn("Zeminde istif (3.1): alıntı olay belgelerinde birebir geçmiyor", sorun)
            self.assertIn("'Kurutma gideri' (15.000 TL) belgelerde doğrulanamadı", sorun)
            self.assertIn("Sonuç 'Kapsamda' ama uygulanabilir veya belirsiz istisna var", sorun)
            self.assertEqual(y["dogrulanamayan_sayilar"], ["95000"])
            md = s["md"].read_text(encoding="utf-8")
            self.assertIn("Örnek Kırtasiye Ticaret Ltd. Şti. deposunda", md)
            self.assertIn("| Ön hesap tazminat | 32.500 |", md)
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Özet", "Teminat ve İstisnalar", "Ön Hesap", "Kontroller"])
            self.assertEqual(wb["Teminat ve İstisnalar"]["C4"].value, "Uygulanır")

    def test_klasor_yok(self):
        with self.assertRaises(llm.LLMHatasi):
            agent.calistir(Path("olmayan_klasor"), Path("x.xlsx"), evet=True)


if __name__ == "__main__":
    unittest.main()
