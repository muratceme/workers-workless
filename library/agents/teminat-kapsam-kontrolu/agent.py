"""
Teminat Kapsam Kontrolü — Workers / Workless AI Agent
Sigorta › Hasar Yönetimi › Hasar Uzmanı

1. Klasördeki dosyalar okunur: hasar bilgileri (.txt: Poliçe Başlangıç/Bitiş, Hasar Tarihi, İhbar Tarihi, Prim
   Durumu, Talep Edilen Teminat, Hasar Tutarı, Sigorta Bedeli, Gerçek Değer), teminat tablosu (Teminat, Limit,
   Muafiyet), poliçe / özel şart / genel şart metinleri ve olay belgeleri (ihbar, eksper raporu, tutanak).
2. Kod kontrol eder ve hesaplar: hasar tarihi poliçe süresinde mi, poliçe başlangıcına yakın mı, prim durumu,
   talep edilen teminat poliçede var mı, ihbar süresi (gün), eksik sigorta oranı (bedel / gerçek değer), muafiyet
   ("%2, en az 5.000 TL" gibi) ve limit ile tazminat ön hesabı.
3. Sigortalı, poliçe no ve --gizle ile verilen kişi adları takma adlarla; telefon, e-posta, TCKN ve IBAN maskelenir.
   Model olayı poliçe teminatları, istisnalar ve şartlarla karşılaştırır; her değerlendirmeyi poliçeden ve olay
   belgelerinden birebir alıntıyla destekler. Kapsam dışı kalabilecek kalemleri tutarıyla belirtirse kod ikinci bir
   ön hesap yapar.
4. Kod modelin çıktısını denetler: alıntılar kaynaklarda birebir geçiyor mu, kapsam dışı tutarlar belgelerde var mı,
   "Kapsamda" sonucu kod bulgularıyla çelişiyor mu; girdilerde olmayan sayılar işaretlenir.
5. Çıktı: Markdown kapsam değerlendirmesi + Excel (Özet + 'Hasar Uzmanı Kararı', Teminat ve İstisnalar,
   Ön Hesap, Kontroller).

Kullanım:
    python agent.py                                          # örnek: kurgusal işyeri dahili su hasarı
    python agent.py --girdi ./hasar_klasoru --gizle "Görüşülen Kişi"
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from dataclasses import dataclass
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
SIFIR = Decimal(0)
YAKIN_BASLANGIC_GUN = 30
POLICE_ADLARI = ("police", "sart", "genel sart", "ozel sart", "kloz", "teminat metni")
OLAY_ADLARI = ("ihbar", "eksper", "tutanak", "beyan", "olay", "rapor", "ifade")


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def katla(s) -> str:
    s = kucuk(s)
    for a, b in zip("ıöüşğçâîû", "iousgcaiu"):
        s = s.replace(a, b)
    return re.sub(r"[^a-z0-9%]+", " ", s).strip()


def duz(s: str) -> str:
    return " ".join(katla(s).split())


def para(x) -> Decimal | None:
    if x in (None, "", "—", "-"):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = re.sub(r"[^\d.,\-]", "", str(x))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tl(x) -> str:
    return "—" if x is None else f"{x:,.0f}".replace(",", ".")


def yuzde(x) -> str:
    return "—" if x is None else f"%{x * 100:.1f}".replace(".", ",")


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(str(x or "").strip(), f).date()
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = belge.txt_oku(yol)
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def anahtarli(metin: str, *etiketler: str) -> str | None:
    for et in etiketler:
        m = re.search(r"^\s*" + re.escape(et) + r"\s*:\s*(.+)$", metin or "", re.I | re.M)
        if m:
            return m.group(1).strip()
    return None


def muafiyet_coz(metin) -> dict:
    """'%2, en az 5.000 TL' → oran 0,02, en az 5000; '10.000 TL' → sabit 10000; '%10' → oran 0,10."""
    ham = str(metin or "").strip()
    k = kucuk(ham)
    for a, b in zip("ıöüşğçâîû", "iousgcaiu"):
        k = k.replace(a, b)
    bul = lambda desen: (lambda m: para(m.group(1).rstrip(".,")) if m else None)(re.search(desen, k))  # noqa: E731
    oran = bul(r"%\s*([\d.,]+)")
    return {"oran": oran / 100 if oran is not None else None, "en_az": bul(r"(?:en az|asgari)\s*([\d.,]+)"),
            "en_cok": bul(r"(?:en fazla|azami)\s*([\d.,]+)"), "sabit": None if oran is not None else bul(r"([\d][\d.,]*)"), "metin": ham}


def muafiyet_tutari(m: dict, hasar: Decimal) -> Decimal:
    if m["sabit"] is not None:
        return m["sabit"]
    if m["oran"] is None:
        return SIFIR
    t = hasar * m["oran"]
    if m["en_az"] is not None:
        t = max(t, m["en_az"])
    if m["en_cok"] is not None:
        t = min(t, m["en_cok"])
    return t


# ----------------------------------------------------------------------------
# Girdiler
# ----------------------------------------------------------------------------

@dataclass
class Dosya:
    ad: str
    tur: str
    metin: str


def dosyalari_oku(klasor: Path) -> tuple[dict, list[dict], list[Dosya], list[str]]:
    bilgi, teminatlar, metinler, atlanan = {}, [], [], []
    for yol in sorted(klasor.iterdir()):
        ad = katla(yol.stem)
        if yol.suffix.lower() in {".csv", ".xlsx", ".xlsm"}:
            s = tablo_oku(yol)
            bi = next((i for i, r in enumerate(s[:10]) if any(katla(c) == "teminat" for c in r)), None)
            if bi is None:
                atlanan.append(f"{yol.name} (Teminat sütunu yok)")
                continue
            b = [katla(c) for c in s[bi]]
            j = lambda *a: next((b.index(x) for x in a if x in b), None)  # noqa: E731
            it, il, im = j("teminat"), j("limit", "teminat limiti", "sigorta bedeli"), j("muafiyet", "tenzili muafiyet")
            for r in s[bi + 1:]:
                if it < len(r) and r[it]:
                    teminatlar.append({"teminat": str(r[it]).strip(), "limit": para(r[il]) if il is not None and il < len(r) else None,
                                       "muafiyet": muafiyet_coz(r[im] if im is not None and im < len(r) else "")})
            continue
        if yol.suffix.lower() not in belge.DESTEKLENEN:
            atlanan.append(yol.name)
            continue
        try:
            metin = belge.metin_oku(yol)
        except belge.BelgeHatasi as h:
            atlanan.append(f"{yol.name} ({h})")
            continue
        if "hasar bilgi" in ad or anahtarli(metin, "Hasar Tarihi") and anahtarli(metin, "Poliçe Başlangıç"):
            bilgi = {"metin": metin, "dosya": yol.name}
            tur = "Hasar bilgileri"
        elif any(a in ad for a in POLICE_ADLARI):
            tur = "Poliçe / şartlar"
        elif any(a in ad for a in OLAY_ADLARI):
            tur = "Olay belgesi"
        else:
            tur = "Poliçe / şartlar" if "teminat disi" in katla(metin) or "istisna" in katla(metin) else "Olay belgesi"
        metinler.append(Dosya(yol.name, tur, metin))
    return bilgi, teminatlar, metinler, atlanan


# ----------------------------------------------------------------------------
# Kontroller ve ön hesap
# ----------------------------------------------------------------------------

def kontrol_et(bilgi: dict, teminatlar: list[dict], metinler: list[Dosya]) -> dict:
    bulgular = []

    def b(onem, aciklama, kaynak=""):
        bulgular.append((onem, aciklama, kaynak))

    t = bilgi.get("metin", "")
    h = {"dosya_no": anahtarli(t, "Dosya No", "Hasar Dosya No"), "sigortali": anahtarli(t, "Sigortalı"), "police_no": anahtarli(t, "Poliçe No"),
         "baslangic": tarih(anahtarli(t, "Poliçe Başlangıç", "Başlangıç Tarihi")), "bitis": tarih(anahtarli(t, "Poliçe Bitiş", "Bitiş Tarihi")),
         "hasar_tarihi": tarih(anahtarli(t, "Hasar Tarihi")), "ihbar_tarihi": tarih(anahtarli(t, "İhbar Tarihi")),
         "prim": anahtarli(t, "Prim Durumu"), "teminat": anahtarli(t, "Talep Edilen Teminat", "Teminat"),
         "hasar": para(anahtarli(t, "Hasar Tutarı")), "bedel": para(anahtarli(t, "Sigorta Bedeli")), "deger": para(anahtarli(t, "Gerçek Değer"))}
    if not bilgi:
        b("orta", "Hasar bilgileri dosyası yok: tarih ve tutar kontrolleri yapılamadı")
    if h["hasar_tarihi"] and h["baslangic"] and h["bitis"]:
        if not (h["baslangic"] <= h["hasar_tarihi"] <= h["bitis"]):
            b("yüksek", f"Hasar tarihi ({h['hasar_tarihi']:%d.%m.%Y}) poliçe süresi dışında ({h['baslangic']:%d.%m.%Y} – {h['bitis']:%d.%m.%Y})", "hasar bilgileri")
        elif (h["hasar_tarihi"] - h["baslangic"]).days <= YAKIN_BASLANGIC_GUN:
            b("orta", f"Hasar poliçe başlangıcından {(h['hasar_tarihi'] - h['baslangic']).days} gün sonra: başlangıca yakın hasar", "hasar bilgileri")
    elif bilgi:
        b("orta", "Poliçe başlangıç/bitiş veya hasar tarihi eksik: süre kontrolü yapılamadı", "hasar bilgileri")
    if h["prim"] and not any(x in katla(h["prim"]) for x in ("odendi", "tahsil edildi", "odenmis")):
        b("yüksek", f"Prim durumu: {h['prim']} — primin ödenmemesinin sonucunu genel şartlara göre değerlendirin", "hasar bilgileri")
    if h["hasar_tarihi"] and h["ihbar_tarihi"]:
        h["ihbar_gun"] = (h["ihbar_tarihi"] - h["hasar_tarihi"]).days
        if h["ihbar_gun"] < 0:
            b("yüksek", "İhbar tarihi hasar tarihinden önce", "hasar bilgileri")
        else:
            b("bilgi", f"İhbar hasardan {h['ihbar_gun']} gün sonra yapılmış; ihbar süresini genel şartlara göre kontrol edin", "hasar bilgileri")
    tem = None
    if h["teminat"]:
        tem = next((x for x in teminatlar if katla(x["teminat"]) == katla(h["teminat"])), None) or \
            next((x for x in teminatlar if katla(h["teminat"]) in katla(x["teminat"]) or katla(x["teminat"]) in katla(h["teminat"])), None)
        if teminatlar and tem is None:
            b("yüksek", f"Talep edilen teminat '{h['teminat']}' poliçe teminat tablosunda yok", "teminatlar")
    if not teminatlar:
        b("orta", "Teminat tablosu (Teminat, Limit, Muafiyet) verilmedi: limit ve muafiyet hesaplanamadı")
    oran = None
    if h["bedel"] and h["deger"] and h["deger"] > h["bedel"]:
        oran = h["bedel"] / h["deger"]
        b("orta", f"Eksik sigorta: sigorta bedeli {tl(h['bedel'])} TL, gerçek değer {tl(h['deger'])} TL (oran {yuzde(oran)})", "hasar bilgileri")
    if not any(d.tur == "Poliçe / şartlar" for d in metinler):
        b("orta", "Poliçe / özel şart / genel şart metni verilmedi: istisnalar değerlendirilemez")
    hesap = on_hesap(h["hasar"], oran, tem)
    if hesap and tem and tem["limit"] is not None and hesap["ara"] > tem["limit"]:
        b("bilgi", f"Hesaplanan tutar teminat limitini ({tl(tem['limit'])} TL) aşıyor; limitle sınırlandı", "teminatlar")
    return {"bulgular": bulgular, "hasar": h, "teminat": tem, "teminatlar": teminatlar, "oran": oran, "hesap": hesap}


def on_hesap(hasar: Decimal | None, oran: Decimal | None, tem: dict | None, haric: Decimal = SIFIR) -> dict | None:
    """Hasar − kapsam dışı → eksik sigorta oranı → muafiyet → limit. Uygulama sırası poliçeye göre değişebilir."""
    if hasar is None:
        return None
    adim = [("Hasar tutarı (eksper)", hasar)]
    tutar = hasar
    if haric:
        tutar -= haric
        adim.append(("Kapsam dışı kalemler", -haric))
    if oran is not None:
        yeni = (tutar * oran).quantize(Decimal(1), ROUND_HALF_UP)
        adim.append((f"Eksik sigorta oranı ({yuzde(oran)})", yeni - tutar))
        tutar = yeni
    muaf = SIFIR
    if tem:
        muaf = muafiyet_tutari(tem["muafiyet"], hasar - haric).quantize(Decimal(1), ROUND_HALF_UP)
        if muaf:
            adim.append((f"Muafiyet ({tem['muafiyet']['metin']})", -min(muaf, tutar)))
            tutar = max(SIFIR, tutar - muaf)
    ara = tutar
    if tem and tem["limit"] is not None and tutar > tem["limit"]:
        adim.append((f"Limit ({tl(tem['limit'])} TL)", tem["limit"] - tutar))
        tutar = tem["limit"]
    adim.append(("Ön hesap tazminat", tutar))
    return {"adimlar": adim, "sonuc": tutar, "ara": ara}


# ----------------------------------------------------------------------------
# Maskeleme
# ----------------------------------------------------------------------------

def takma_adlar(h: dict, metinler: list[Dosya], terimler: list[str]) -> dict[str, str]:
    harita = {}
    if h.get("sigortali"):
        harita[h["sigortali"]] = "[SİGORTALI]"
        kok = re.split(r"\s+(?:San\.|Sanayi|Tic\.|Ticaret|A\.Ş\.|Ltd\.)", h["sigortali"])[0].strip()
        if len(kok) > 4 and kok != h["sigortali"]:
            harita[kok] = "[SİGORTALI]"
    if h.get("police_no"):
        harita[h["police_no"]] = "[POLİÇE NO]"
    kisiler = list(terimler)
    for d in metinler:
        for m in re.finditer(r"(?:yetkilisi|Görüşülen|Beyan veren|Sürücü)\s*:?\s+([A-ZÇĞİÖŞÜ][a-zçğıöşü]+(?:\s+[A-ZÇĞİÖŞÜ][a-zçğıöşü]+)+?)(?='|’|\s*[,.(:]|\s+(?:ile|tarafından|beyan))",
                             d.metin):
            kisiler.append(m.group(1))
    for i, ad in enumerate(dict.fromkeys(a.strip() for a in kisiler if a.strip()), 1):
        harita.setdefault(ad, f"[KİŞİ-{i}]")
    return harita


def maskele(metin: str, harita: dict[str, str]) -> str:
    for gercek, takma in sorted(harita.items(), key=lambda x: -len(x[0])):
        metin = re.sub(re.escape(gercek), takma, metin, flags=re.I)
    return llm.maskele(metin)


def geri_ac(metin: str, harita: dict[str, str]) -> str:
    for gercek, takma in sorted(harita.items(), key=lambda x: (-len(x[1]), -len(x[0]))):
        metin = metin.replace(takma, gercek)
    return metin


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

ALINTI = {"type": "object", "properties": {"kaynak": {"type": "string"}, "alinti": {"type": "string"}},
          "required": ["kaynak", "alinti"], "additionalProperties": False}
SEMA = {
    "type": "object",
    "properties": {
        "olay_ozeti": {"type": "string"},
        "teminat_degerlendirmesi": {"type": "array", "items": {
            "type": "object",
            "properties": {"teminat": {"type": "string"}, "sonuc": {"type": "string", "enum": ["Kapsamda", "Kapsam dışı", "Belirsiz"]},
                           "gerekce": {"type": "string"}, "police_alintilari": {"type": "array", "items": ALINTI},
                           "olay_alintilari": {"type": "array", "items": ALINTI}},
            "required": ["teminat", "sonuc", "gerekce", "police_alintilari", "olay_alintilari"], "additionalProperties": False}},
        "istisnalar": {"type": "array", "items": {
            "type": "object",
            "properties": {"istisna": {"type": "string"}, "uygulanir": {"type": "string", "enum": ["Uygulanır", "Uygulanmaz", "Belirsiz"]},
                           "gerekce": {"type": "string"}, "police_alintilari": {"type": "array", "items": ALINTI},
                           "olay_alintilari": {"type": "array", "items": ALINTI}},
            "required": ["istisna", "uygulanir", "gerekce", "police_alintilari", "olay_alintilari"], "additionalProperties": False}},
        "kapsam_disi_kalemler": {"type": "array", "items": {
            "type": "object", "properties": {"aciklama": {"type": "string"}, "tutar": {"type": "number"}, "olay_alintisi": {"type": "string"}},
            "required": ["aciklama", "tutar", "olay_alintisi"], "additionalProperties": False}},
        "sonuc": {"type": "object", "properties": {
            "degerlendirme": {"type": "string", "enum": ["Kapsamda", "Kısmen kapsamda", "Kapsam dışı", "Ek bilgi gerekli"]},
            "gerekce": {"type": "string"}}, "required": ["degerlendirme", "gerekce"], "additionalProperties": False},
        "ek_bilgi_gerekenler": {"type": "array", "items": {"type": "string"}},
        "celiskiler": {"type": "array", "items": {"type": "string"}},
    },
    "required": ["olay_ozeti", "teminat_degerlendirmesi", "istisnalar", "kapsam_disi_kalemler", "sonuc", "ek_bilgi_gerekenler", "celiskiler"],
    "additionalProperties": False,
}


def sayilar(metin: str) -> set[str]:
    return {re.sub(r"\D", "", m) for m in re.findall(r"\d[\d.,]*\d", metin) if len(re.sub(r"\D", "", m)) >= 3}


def kod_ozeti(k: dict) -> list[str]:
    h, s = k["hasar"], []
    for et, a in (("Poliçe başlangıç", "baslangic"), ("Poliçe bitiş", "bitis"), ("Hasar tarihi", "hasar_tarihi"), ("İhbar tarihi", "ihbar_tarihi")):
        if h[a]:
            s.append(f"- {et}: {h[a]:%d.%m.%Y}")
    for et, a in (("Talep edilen teminat", "teminat"), ("Prim durumu", "prim")):
        if h[a]:
            s.append(f"- {et}: {h[a]}")
    for et, a in (("Hasar tutarı", "hasar"), ("Sigorta bedeli", "bedel"), ("Gerçek değer", "deger")):
        if h[a] is not None:
            s.append(f"- {et}: {tl(h[a])} TL")
    s += [f"- Teminat: {x['teminat']}, limit {tl(x['limit'])} TL, muafiyet {x['muafiyet']['metin'] or 'yok'}" for x in k["teminatlar"]]
    if k["hesap"]:
        s += [f"- Ön hesap adımı: {a} {tl(v)} TL" for a, v in k["hesap"]["adimlar"]]
    return s


def hazirla(metinler, k, harita) -> tuple[str, str]:
    mesaj = "\n".join([
        "<police_metinleri>", *[f'<belge ad="{d.ad}">\n{maskele(d.metin, harita)}\n</belge>' for d in metinler if d.tur == "Poliçe / şartlar"],
        "</police_metinleri>",
        "<olay_belgeleri>", *[f'<belge ad="{d.ad}">\n{maskele(d.metin, harita)}\n</belge>' for d in metinler if d.tur != "Poliçe / şartlar"],
        "</olay_belgeleri>",
        "<kod_ozeti>", *[maskele(x, harita) for x in kod_ozeti(k)], "</kod_ozeti>",
        "<kod_kontrolleri>", *([f"- [{o}] {maskele(a, harita)}" for o, a, _ in k["bulgular"]] or ["- Kontrol bulgusu yok"]), "</kod_kontrolleri>",
    ])
    return (BURASI / "prompt.md").read_text(encoding="utf-8"), mesaj


def denetle(y: dict, k: dict, metinler: list[Dosya], harita: dict, mesaj: str) -> list[tuple[str, str, str]]:
    sorunlar = []
    police = duz(" ".join(maskele(d.metin, harita) for d in metinler if d.tur == "Poliçe / şartlar"))
    olay = duz(" ".join(maskele(d.metin, harita) for d in metinler if d.tur != "Poliçe / şartlar"))
    for grup, ad in (("teminat_degerlendirmesi", "teminat"), ("istisnalar", "istisna")):
        for x in y[grup]:
            for alan, kaynak, et in (("police_alintilari", police, "poliçede"), ("olay_alintilari", olay, "olay belgelerinde")):
                for a in x[alan]:
                    a["dogrulandi"] = bool(a["alinti"].strip()) and duz(a["alinti"]) in kaynak
                    if not a["dogrulandi"]:
                        sorunlar.append(("yüksek", f"{x[ad]}: alıntı {et} birebir geçmiyor — \"{a['alinti'][:90]}\"", "model ↔ belge"))
            if grup == "istisnalar" and not x["police_alintilari"]:
                sorunlar.append(("orta", f"İstisna '{x['istisna']}' için poliçe alıntısı yok", "model"))
    haric = SIFIR
    bilinen_tutarlar = sayilar(mesaj)
    for x in y["kapsam_disi_kalemler"]:
        tutar = Decimal(str(x["tutar"])).quantize(Decimal(1), ROUND_HALF_UP)
        x["dogrulandi"] = str(tutar) in bilinen_tutarlar and bool(x["olay_alintisi"].strip()) and duz(x["olay_alintisi"]) in olay
        if x["dogrulandi"]:
            haric += tutar
        else:
            sorunlar.append(("yüksek", f"Kapsam dışı kalem '{x['aciklama']}' ({tl(tutar)} TL) belgelerde doğrulanamadı; ön hesaba alınmadı", "model ↔ belge"))
    y["ikinci_hesap"] = on_hesap(k["hasar"]["hasar"], k["oran"], k["teminat"], haric) if haric else None
    sonuc = y["sonuc"]["degerlendirme"]
    yuksek = [a for o, a, _ in k["bulgular"] if o == "yüksek"]
    if sonuc == "Kapsamda" and yuksek:
        sorunlar.append(("yüksek", f"Sonuç 'Kapsamda' ama kodda yüksek önemli bulgu var: {yuksek[0][:100]}", "model ↔ kod"))
    if sonuc == "Kapsamda" and any(x["uygulanir"] != "Uygulanmaz" for x in y["istisnalar"]):
        sorunlar.append(("orta", "Sonuç 'Kapsamda' ama uygulanabilir veya belirsiz istisna var", "model"))
    bilinen = bilinen_tutarlar | {x for x in re.findall(r"\d+", mesaj) if len(x) >= 3}
    metin = " ".join([y["olay_ozeti"], y["sonuc"]["gerekce"], *y["ek_bilgi_gerekenler"], *y["celiskiler"],
                      *[x["gerekce"] for g in ("teminat_degerlendirmesi", "istisnalar") for x in y[g]], *[x["aciklama"] for x in y["kapsam_disi_kalemler"]]])
    y["dogrulanamayan_sayilar"] = sorted(s for s in sayilar(metin) if s not in bilinen)
    return sorunlar


def model_yaz(metinler, k, harita) -> dict:
    sistem, mesaj = hazirla(metinler, k, harita)
    y = llm.json_iste(sistem, mesaj, SEMA)
    y["sorunlar"] = denetle(y, k, metinler, harita, mesaj)
    return y


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"yüksek": "FDE2E1", "orta": "FFF4CE", "bilgi": "E8F0FE", "Kapsamda": "E3F4E1", "Kapsam dışı": "FDE2E1", "Belirsiz": "FFF4CE",
        "Uygulanır": "FDE2E1", "Uygulanmaz": "E3F4E1"}
KARAR = PatternFill("solid", fgColor="FFF4CE")
MODEL = PatternFill("solid", fgColor="EDE7FF")
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def _alintilar(lst, g) -> str:
    return "\n".join(f"[{a['kaynak']}] \"{g(a['alinti'])}\"" + ("" if a["dogrulandi"] else " [DOĞRULANAMADI]") for a in lst)


def rapor_yaz(cikti: Path, metinler, atlanan, k: dict, y: dict, harita: dict) -> Path:
    g = lambda s: geri_ac(s, harita)  # noqa: E731
    h, tum = k["hasar"], k["bulgular"] + y["sorunlar"]
    wb = Workbook()
    o = wb.active
    o.title = "Özet"
    o.column_dimensions["A"].width, o.column_dimensions["B"].width = 30, 100
    for a, d in [("Dosya No", h["dosya_no"] or "—"), ("Sigortalı", h["sigortali"] or "—"), ("Poliçe No", h["police_no"] or "—"),
                 ("Talep edilen teminat", h["teminat"] or "—"), ("Kaynaklar", "; ".join(f"{d.ad} ({d.tur})" for d in metinler)),
                 ("Okunamayan", "; ".join(atlanan) or "—"), ("Olay özeti (model)", g(y["olay_ozeti"])),
                 ("Kapsam değerlendirmesi (taslak, model)", f"{y['sonuc']['degerlendirme']} — {g(y['sonuc']['gerekce'])}"),
                 ("Ön hesap (kod)", f"{tl(k['hesap']['sonuc'])} TL" if k["hesap"] else "—"),
                 ("Kapsam dışı kalemler düşülerek (kod)", f"{tl(y['ikinci_hesap']['sonuc'])} TL" if y["ikinci_hesap"] else "—"),
                 ("Çelişkiler (model)", "\n".join(f"- {g(x)}" for x in y["celiskiler"]) or "—"),
                 ("Ek bilgi gerekenler (model)", "\n".join(f"- {g(x)}" for x in y["ek_bilgi_gerekenler"]) or "—"),
                 ("Doğrulanamayan sayılar", ", ".join(y["dogrulanamayan_sayilar"]) or "—"), ("Model", llm.kullanim_ozeti()),
                 ("Hasar Uzmanı Kararı", ""), ("Karar Notu", "")]:
        o.append([a, d])
        o.cell(o.max_row, 1).font = Font(bold=True)
        if "model" in a:
            o.cell(o.max_row, 2).fill = MODEL
        if a.startswith(("Hasar Uzmanı", "Karar")):
            o.cell(o.max_row, 2).fill = KARAR
        if a == "Doğrulanamayan sayılar" and y["dogrulanamayan_sayilar"]:
            o.cell(o.max_row, 2).fill = PatternFill("solid", fgColor=RENK["yüksek"])
        for c in o[o.max_row]:
            c.alignment = UST

    ti = wb.create_sheet("Teminat ve İstisnalar")
    _baslik(ti, ["Tür", "Konu", "Sonuç", "Gerekçe", "Poliçe Alıntıları", "Olay Alıntıları", "İnceleme"], (10, 30, 13, 50, 60, 60, 24))
    for x in y["teminat_degerlendirmesi"]:
        ti.append(["Teminat", g(x["teminat"]), x["sonuc"], g(x["gerekce"]), _alintilar(x["police_alintilari"], g), _alintilar(x["olay_alintilari"], g), ""])
        ti.cell(ti.max_row, 3).fill = PatternFill("solid", fgColor=RENK[x["sonuc"]])
    for x in y["istisnalar"]:
        ti.append(["İstisna", g(x["istisna"]), x["uygulanir"], g(x["gerekce"]), _alintilar(x["police_alintilari"], g), _alintilar(x["olay_alintilari"], g), ""])
        ti.cell(ti.max_row, 3).fill = PatternFill("solid", fgColor=RENK[x["uygulanir"]])
    for r in ti.iter_rows(min_row=2):
        r[6].fill = KARAR
        for c in r:
            c.alignment = UST

    oh = wb.create_sheet("Ön Hesap")
    _baslik(oh, ["Senaryo", "Adım", "Tutar (TL)"], (34, 50, 16))
    for ad, hs in (("1 — Tüm hasar", k["hesap"]), ("2 — Kapsam dışı kalemler düşülerek", y["ikinci_hesap"])):
        if hs:
            for a, v in hs["adimlar"]:
                oh.append([ad, a, float(v)])
                oh.cell(oh.max_row, 3).number_format = "#,##0"
                if a == "Ön hesap tazminat":
                    for c in oh[oh.max_row]:
                        c.font = Font(bold=True)
    for x in y["kapsam_disi_kalemler"]:
        oh.append(["Kapsam dışı kalem (model)", g(x["aciklama"]) + ("" if x["dogrulandi"] else " [DOĞRULANAMADI, hesaba alınmadı]"), x["tutar"]])
        oh.cell(oh.max_row, 3).number_format = "#,##0"
    oh.append([])
    oh.append(["Not", "Sıra: kapsam dışı → eksik sigorta oranı → muafiyet → limit. Uygulama sırası ve değer esası poliçenize göre değişebilir."])

    kt = wb.create_sheet("Kontroller")
    _baslik(kt, ["Önem", "Kontrol", "Kaynak", "İnceleme"], (9, 100, 18, 24))
    for on, a, kay in tum:
        kt.append([on, g(a), kay, ""])
        kt.cell(kt.max_row, 1).fill = PatternFill("solid", fgColor=RENK[on])
        kt.cell(kt.max_row, 4).fill = KARAR
        kt.cell(kt.max_row, 2).alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

    md = cikti.with_suffix(".md")
    sat = [f"# Teminat Kapsam Değerlendirmesi (taslak) — {h['dosya_no'] or ''}", "",
           f"Sigortalı: {h['sigortali'] or '—'} · Talep edilen teminat: {h['teminat'] or '—'}", "", "## Olay", "", g(y["olay_ozeti"]), "",
           "## Teminat", ""]
    for x in y["teminat_degerlendirmesi"]:
        sat += [f"### {g(x['teminat'])} — {x['sonuc']}", "", g(x["gerekce"]), ""]
        sat += [f"- Poliçe [{a['kaynak']}]: \"{g(a['alinti'])}\"" + ("" if a["dogrulandi"] else " **[doğrulanamadı]**") for a in x["police_alintilari"]]
        sat += [f"- Olay [{a['kaynak']}]: \"{g(a['alinti'])}\"" + ("" if a["dogrulandi"] else " **[doğrulanamadı]**") for a in x["olay_alintilari"]]
        sat += [""]
    sat += ["## İstisnalar ve şartlar", ""]
    for x in y["istisnalar"]:
        sat += [f"### {g(x['istisna'])} — {x['uygulanir']}", "", g(x["gerekce"]), ""]
        sat += [f"- Poliçe [{a['kaynak']}]: \"{g(a['alinti'])}\"" + ("" if a["dogrulandi"] else " **[doğrulanamadı]**") for a in x["police_alintilari"]]
        sat += [f"- Olay [{a['kaynak']}]: \"{g(a['alinti'])}\"" + ("" if a["dogrulandi"] else " **[doğrulanamadı]**") for a in x["olay_alintilari"]]
        sat += [""]
    sat += ["## Ön hesap (kod)", ""]
    for ad, hs in (("Tüm hasar", k["hesap"]), ("Kapsam dışı kalemler düşülerek", y["ikinci_hesap"])):
        if hs:
            sat += [f"**{ad}**", "", "| Adım | Tutar (TL) |", "|---|---:|"] + [f"| {a} | {tl(v)} |" for a, v in hs["adimlar"]] + [""]
    sat += ["## Sonuç (taslak)", "", f"**{y['sonuc']['degerlendirme']}** — {g(y['sonuc']['gerekce'])}", ""]
    if y["celiskiler"]:
        sat += ["## Çelişkiler", ""] + [f"- {g(x)}" for x in y["celiskiler"]] + [""]
    if y["ek_bilgi_gerekenler"]:
        sat += ["## Ek bilgi / belge", ""] + [f"- {g(x)}" for x in y["ek_bilgi_gerekenler"]] + [""]
    sat += ["## Kontroller", ""] + ([f"- **{o_}** {g(a)}" for o_, a, _ in tum] or ["- Bulgu yok"])
    if y["dogrulanamayan_sayilar"]:
        sat += ["", f"> Girdilerde bulunmayan sayılar (kontrol edin): {', '.join(y['dogrulanamayan_sayilar'])}"]
    sat += ["", "> Taslaktır; kapsam ve tazminat kararı yetkili hasar uzmanına aittir. Poliçe genel ve özel şartlarının tamamını esas alın."]
    md.write_text("\n".join(sat) + "\n", encoding="utf-8")
    return md


def calistir(klasor: Path, cikti: Path, terimler: list[str] | None = None, evet: bool = False) -> dict:
    if not klasor.is_dir():
        raise llm.LLMHatasi(f"{klasor} klasörü bulunamadı.")
    bilgi, teminatlar, metinler, atlanan = dosyalari_oku(klasor)
    if not metinler:
        raise llm.LLMHatasi(f"{klasor}: okunabilir belge yok (.txt/.pdf/.docx).")
    k = kontrol_et(bilgi, teminatlar, metinler)
    harita = takma_adlar(k["hasar"], metinler, terimler or [])
    for d in metinler:
        print(f"[OK] {d.ad}: {d.tur}")
    llm.onay_al(f"{len(metinler)} belgenin metni gönderilecek (sigortalı, poliçe no ve {sum(t.startswith('[KİŞİ') for t in harita.values())} kişi "
                "takma adlı; telefon, e-posta, TCKN, IBAN maskeli).", evet)
    y = model_yaz(metinler, k, harita)
    md = rapor_yaz(cikti, metinler, atlanan, k, y, harita)
    return {"kontrol": k, "yanit": y, "harita": harita, "md": md, "metinler": metinler}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Hasar olayını poliçe teminatları, muafiyetler ve istisnalarla karşılaştırıp kapsam değerlendirmesi taslağı hazırlar.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "ornek_dahili_su", help="Hasar klasörü")
    p.add_argument("--gizle", nargs="*", default=[], metavar="AD", help="Ayrıca maskelenecek adlar")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "teminat_kapsam.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.cikti, a.gizle, a.evet)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    for o, ac, _ in s["kontrol"]["bulgular"] + s["yanit"]["sorunlar"]:
        if o != "bilgi":
            print(f"[{'X' if o == 'yüksek' else '!'}] {ac}")
    if s["yanit"]["dogrulanamayan_sayilar"]:
        print(f"[!] Girdilerde olmayan sayılar: {', '.join(s['yanit']['dogrulanamayan_sayilar'])}")
    print(f"[i] Kapsam (taslak): {s['yanit']['sonuc']['degerlendirme']}" +
          (f" · ön hesap {tl(s['kontrol']['hesap']['sonuc'])} TL" if s["kontrol"]["hesap"] else ""))
    print(f"[OK] Rapor: {a.cikti.resolve()} · {s['md'].resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
