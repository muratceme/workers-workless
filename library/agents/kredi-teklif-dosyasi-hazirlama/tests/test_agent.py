"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

DOSYA = KLASOR / "ornek_veri" / "teklif_ornek_ambalaj"
GIZLI = ["Kemal Örnekoğlu", "Selin Örnekoğlu"]
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return {"firma_ve_talep": "[FİRMA] ambalaj üreticisidir; ortakları [GİZLİ-2] ve [GİZLİ-1].",
            "mali_degerlendirme": "Net satışlar %24,2 büyümüş; net finansal borç/FAVÖK 1,90.",
            "istihbarat_degerlendirmesi": "Bir adet ödenmiş karşılıksız çek kaydı var.",
            "teminat_degerlendirmesi": "Karşılama oranı %90,0; ek teminat 2.750.000 TL önerilir.",
            "guclu_yonler": ["Ciro büyümesi güçlü."], "riskler": [{"aciklama": "Limit doluluğu yüksek.", "kaynak": "kod hesapları"}],
            "onerilen_sartlar": ["Çek akışı şartı."], "eksik_bilgiler": ["2026 ara dönem mizanı."],
            "teklif_ozeti": "[FİRMA] için talep, ek teminat koşuluyla değerlendirilebilir."}


class KontrolTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.girdi, cls.metinler, cls.atlanan = agent.dosyalari_oku(DOSYA)
        cls.k = agent.kontrol_et(cls.girdi, cls.metinler, agent.VARSAYILAN_KATSAYI)

    def test_dosya_turleri(self):
        self.assertEqual({d.ad: d.tur for d in self.metinler}, {
            "istihbarat_notlari.txt": "İstihbarat", "kkb_risk.csv": "KKB / memzuç risk", "mali_tablolar.csv": "Mali tablolar",
            "talep_ve_firma_bilgisi.txt": "Talep ve firma bilgisi", "teminatlar.csv": "Teminatlar"})
        self.assertEqual(self.atlanan, [])

    def test_talep(self):
        self.assertEqual(self.k["talep"], {"unvan": "Örnek Ambalaj Sanayi ve Ticaret A.Ş.", "vkn": "1234567890", "limit": D("20000000")})

    def test_rasyolar_elle(self):
        r = self.k["rasyolar"]["2025"]
        self.assertEqual(r["favok"], D("15500000"))                         # 25 − 12,5 + 3 milyon
        self.assertEqual(r["netborc"], D("29500000"))                       # 24 + 8 − 2,5
        self.assertAlmostEqual(float(r["cari"]), 55 / 44, places=6)
        self.assertAlmostEqual(float(r["nb_favok"]), 29.5 / 15.5, places=6)
        self.assertAlmostEqual(float(r["faiz_karsilama"]), 15.5 / 7.5, places=6)
        self.assertAlmostEqual(float(r["alacak_gun"]), 26 * 365 / 118, places=6)
        self.assertEqual(r["denklik"], 0)
        self.assertEqual(self.k["rasyolar"]["2024"]["favok"], D("13500000"))

    def test_bulgular(self):
        metin = "\n".join(b[1] for b in self.k["bulgular"])
        self.assertIn("finansman giderleri %87,5 arttı", metin)
        self.assertIn("nakdi limit doluluğu %95,1", metin)
        self.assertIn("teminat karşılama oranı %90,0", metin)              # (20 × 0,75 + 6 × 0,5) / 20 milyon
        self.assertIn("'karşılıksız'", metin)
        self.assertIn("'gecikme'", metin)
        self.assertNotIn("protest", metin)                                   # "bulunmamaktadır" olumsuz sayılmaz
        self.assertEqual(self.k["teminat_toplam"], D("18000000"))

    def test_denk_olmayan_bilanco(self):
        yillar, veri = self.girdi["mali"]
        bozuk = {y: dict(v) for y, v in veri.items()}
        bozuk["2025"]["ozk"] = D("-1000000")
        k = agent.kontrol_et({**self.girdi, "mali": (yillar, bozuk)}, self.metinler, agent.VARSAYILAN_KATSAYI)
        metin = "\n".join(b[1] for b in k["bulgular"])
        self.assertIn("denk değil", metin)
        self.assertIn("özkaynak negatif", metin)


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_uctan_uca(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "k.xlsx"
            s = agent.calistir(DOSYA, cikti, GIZLI, evet=True)
            m = MESAJLAR[0]
            for gizli in ("Örnek Ambalaj", "1234567890", "Kemal Örnekoğlu", "Selin", "Örnekoğlu", "0532 111 22 33"):
                self.assertNotIn(gizli, m)
            for var in ("[FİRMA]", "<kod_hesaplari>", "<kod_kontrolleri>", "net finansal borç/FAVÖK 1,90"):
                self.assertIn(var, m)
            self.assertEqual(s["yanit"]["dogrulanamayan_sayilar"], ["2750000"])   # uydurma ek teminat tutarı
            md = s["md"].read_text(encoding="utf-8")
            self.assertIn("Örnek Ambalaj Sanayi ve Ticaret A.Ş. ambalaj", md)      # rapor şirket içi: takma adlar geri açılır
            self.assertIn("Kemal Örnekoğlu", md)
            self.assertNotIn("[FİRMA]", md)
            self.assertIn("| Net finansal borç / FAVÖK | 1,56 | 1,90 |", md)
            self.assertIn("2750000", md)
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Özet", "Kontroller", "Rasyolar", "KKB Risk", "Teminatlar"])
            self.assertEqual(wb["Kontroller"].cell(1, 5).value, "Portföy Yöneticisi Onayı")

    def test_klasor_yok(self):
        with self.assertRaises(llm.LLMHatasi):
            agent.calistir(Path("olmayan_klasor"), Path("x.xlsx"), evet=True)


class YardimciTest(unittest.TestCase):
    def test_sayilar(self):
        self.assertEqual(agent.sayilar("15.000.000 TL ve %61,4, 12 ay"), {"15000000", "614"})

    def test_katsayi_dosyasi(self):
        with tempfile.TemporaryDirectory() as t:
            (Path(t) / "k.csv").write_text("Tür;Katsayı\nİpotek;60\nMüşteri Çeki;0,4\n", encoding="utf-8")
            k = agent.katsayi_oku(Path(t) / "k.csv")
        self.assertEqual(k, {"ipotek": D("0.6"), "musteri ceki": D("0.4")})


if __name__ == "__main__":
    unittest.main()
