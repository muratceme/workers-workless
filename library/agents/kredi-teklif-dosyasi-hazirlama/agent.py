"""
Kredi Teklif Dosyası Hazırlama — Workers / Workless AI Agent
Bankacılık › Ticari Bankacılık › Ticari Portföy Yöneticisi

1. Teklif klasöründeki dosyalar okunur: talep bilgisi (.txt), mali tablolar (Kalem × yıl, .csv/.xlsx), KKB/memzuç
   risk dağılımı (.csv/.xlsx), teminatlar (.csv/.xlsx), istihbarat ve diğer notlar (.txt/.pdf/.docx).
2. Kod hesaplar ve kontrol eder: rasyolar (cari oran, asit-test, kaldıraç, net finansal borç/FAVÖK, faiz karşılama,
   alacak/stok devir süreleri, büyüme), bilanço denkliği, KKB riski ↔ bilançodaki banka kredileri, sektördeki limit
   doluluğu, teminat karşılama oranı (teminat türüne göre katsayılı) ve istihbarattaki olumsuz ifadeler.
3. Firma unvanı, VKN, --gizle ile verilen kişi adları; telefon, e-posta, TCKN ve IBAN maskelenir. Model kod
   bulgularını esas alarak tahsis birimi için teklif özeti yazar; modelin metnindeki, girdilerde olmayan sayılar kodla
   işaretlenir.
4. Çıktı: Markdown teklif özeti (tablolar koddan) + Excel (Özet, Kontroller + 'Portföy Yöneticisi Onayı', Rasyolar,
   KKB Risk, Teminatlar).

Kullanım:
    python agent.py                                          # örnek teklif klasörüyle dener
    python agent.py --girdi ./teklif_klasoru --gizle "Ortak Adı" "Kefil Adı"
    python agent.py --girdi ./teklif --katsayilar katsayilar.csv
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from dataclasses import dataclass
from datetime import datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
SIFIR = Decimal(0)

# Teminat değer katsayıları — ÖRNEKTİR; bankanızın teminat politikasına göre --katsayilar ile değiştirin.
VARSAYILAN_KATSAYI = {"ipotek": Decimal("0.75"), "nakit": Decimal("1"), "mevduat": Decimal("1"), "blokaj": Decimal("1"),
                      "kgf": Decimal("1"), "teminat mektubu": Decimal("1"), "musteri ceki": Decimal("0.5"), "cek": Decimal("0.5"),
                      "senet": Decimal("0.3"), "arac rehni": Decimal("0.5"), "tasit rehni": Decimal("0.5"), "makine rehni": Decimal("0.4"),
                      "ticari isletme rehni": Decimal("0.3"), "alacak temliki": Decimal("0.4"), "temlik": Decimal("0.4"),
                      "kefalet": Decimal("0"), "sahsi kefalet": Decimal("0")}
OLUMSUZ = [("karşılıksız", "yüksek"), ("protest", "yüksek"), ("haciz", "yüksek"), ("icra", "yüksek"), ("iflas", "yüksek"),
           ("konkordato", "yüksek"), ("takibe", "yüksek"), ("yasal takip", "yüksek"), ("gecikme", "orta"), ("yapılandır", "orta"),
           ("olumsuz", "orta"), ("şikâyet", "orta"), ("dava", "orta"), ("geç ödeme", "orta")]


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def katla(s) -> str:
    s = kucuk(s)
    for a, b in zip("ıöüşğçâîû", "iousgcaiu"):
        s = s.replace(a, b)
    return re.sub(r"[^a-z0-9%]+", " ", s).strip()


def para(x) -> Decimal | None:
    if x in (None, "", "—", "-"):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("₺", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tl(x, kesir=0) -> str:
    if x is None:
        return "—"
    return f"{x:,.{kesir}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def oran(x, kesir=2) -> str:
    return "—" if x is None else f"{x:.{kesir}f}".replace(".", ",")


def yuzde(x) -> str:
    return "—" if x is None else f"%{x * 100:.1f}".replace(".", ",")


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


# ----------------------------------------------------------------------------
# Girdiler
# ----------------------------------------------------------------------------

@dataclass
class Dosya:
    ad: str
    tur: str
    metin: str


MALI = {  # kod adı → başlık eşanlamlıları (katlanmış)
    "dv": ("donen varliklar",), "hazir": ("hazir degerler", "nakit ve nakit benzerleri"), "alacak": ("ticari alacaklar",),
    "stok": ("stoklar",), "duran": ("duran varliklar",), "aktif": ("toplam aktif", "aktif toplami", "toplam varliklar"),
    "kvyk": ("kisa vadeli yabanci kaynaklar", "kisa vadeli yukumlulukler"), "kvbk": ("kisa vadeli banka kredileri", "kv banka kredileri", "kisa vadeli mali borclar"),
    "uvyk": ("uzun vadeli yabanci kaynaklar", "uzun vadeli yukumlulukler"), "uvbk": ("uzun vadeli banka kredileri", "uv banka kredileri", "uzun vadeli mali borclar"),
    "ozk": ("ozkaynaklar", "ozsermaye"), "satis": ("net satislar", "hasilat"), "smm": ("satislarin maliyeti", "satilan mal maliyeti"),
    "brut": ("brut kar", "brut satis kari"), "fgid": ("faaliyet giderleri",), "amort": ("amortisman", "amortisman ve itfa"),
    "favok": ("favok", "ebitda"), "fin": ("finansman giderleri", "finansman gideri"), "net": ("net kar", "donem net kari", "net donem kari"),
}


def mali_oku(yol: Path) -> tuple[list[str], dict[str, dict[str, Decimal | None]]]:
    s = tablo_oku(yol)
    bi = next((i for i, r in enumerate(s[:10]) if katla(r[0]) in ("kalem", "hesap", "kalem adi")), 0)
    yillar = [str(c).strip().removesuffix(".0") for c in s[bi][1:] if c not in (None, "")]
    veri: dict[str, dict] = {y: {} for y in yillar}
    for r in s[bi + 1:]:
        k = katla(r[0])
        kod = next((kd for kd, adlar in MALI.items() if k in adlar), None)
        if kod:
            for j, y in enumerate(yillar, 1):
                veri[y][kod] = para(r[j]) if j < len(r) else None
    return yillar, veri


def kkb_oku(yol: Path) -> list[dict]:
    s = tablo_oku(yol)
    bi = next((i for i, r in enumerate(s[:10]) if any(katla(c) in ("banka", "kurum") for c in r)), 0)
    b = [katla(c) for c in s[bi]]

    def i(*adlar):
        return next((b.index(a) for a in adlar if a in b), None)
    ib, inl, inr, igl, igr = (i("banka", "kurum"), i("nakdi limit"), i("nakdi risk"), i("gayrinakdi limit", "gayri nakdi limit"),
                              i("gayrinakdi risk", "gayri nakdi risk"))
    sonuc = []
    for r in s[bi + 1:]:
        if not r[ib] or katla(r[ib]).startswith(("toplam", "genel")):
            continue
        al = lambda j: (para(r[j]) or SIFIR) if j is not None and j < len(r) else SIFIR  # noqa: E731
        sonuc.append({"banka": str(r[ib]).strip(), "nl": al(inl), "nr": al(inr), "gl": al(igl), "gr": al(igr)})
    return sonuc


def teminat_oku(yol: Path) -> list[dict]:
    s = tablo_oku(yol)
    bi = next((i for i, r in enumerate(s[:10]) if any(katla(c) in ("tur", "teminat turu", "teminat") for c in r)), 0)
    b = [katla(c) for c in s[bi]]
    it = next(b.index(a) for a in ("teminat turu", "tur", "teminat") if a in b)
    ia = next((b.index(a) for a in ("aciklama", "detay") if a in b), None)
    iv = next((b.index(a) for a in ("deger", "ekspertiz degeri", "tutar", "nominal") if a in b), None)
    return [{"tur": str(r[it]).strip(), "aciklama": str(r[ia] or "").strip() if ia is not None and ia < len(r) else "",
             "deger": para(r[iv]) if iv is not None and iv < len(r) else None} for r in s[bi + 1:] if r[it]]


def dosyalari_oku(klasor: Path) -> tuple[dict, list[Dosya], list[str]]:
    girdi = {"talep": None, "mali": None, "kkb": None, "teminat": None}
    metinler, atlanan = [], []
    for yol in sorted(klasor.iterdir()):
        ad = katla(yol.stem)
        if yol.suffix.lower() in {".csv", ".xlsx", ".xlsm"}:
            ilk = katla(" ".join(str(c) for r in tablo_oku(yol)[:3] for c in r))
            if "mali" in ad or "bilanco" in ad or "kalem" in ilk:
                girdi["mali"], tur = mali_oku(yol), "Mali tablolar"
            elif "kkb" in ad or "memzuc" in ad or "risk" in ad or "nakdi" in ilk:
                girdi["kkb"], tur = kkb_oku(yol), "KKB / memzuç risk"
            elif "teminat" in ad or "teminat" in ilk:
                girdi["teminat"], tur = teminat_oku(yol), "Teminatlar"
            else:
                atlanan.append(f"{yol.name} (tanınmayan tablo)")
                continue
            metinler.append(Dosya(yol.name, tur, ""))         # tablolar modele kod hesabı olarak gider
            continue
        if yol.suffix.lower() not in belge.DESTEKLENEN:
            atlanan.append(yol.name)
            continue
        try:
            metin = belge.metin_oku(yol)
        except belge.BelgeHatasi as h:
            atlanan.append(f"{yol.name} ({h})")
            continue
        if "talep" in ad or "basvuru" in ad:
            tur = "Talep ve firma bilgisi"
            girdi["talep"] = metin
        elif "istihbarat" in ad or "istihbarat" in katla(metin[:200]):
            tur = "İstihbarat"
        else:
            tur = "Diğer not"
        metinler.append(Dosya(yol.name, tur, metin))
    return girdi, metinler, atlanan


def alan(metin: str, *etiketler: str) -> str | None:
    for et in etiketler:
        m = re.search(r"^\s*" + re.escape(et) + r"\s*:\s*(.+)$", metin, re.I | re.M)
        if m:
            return m.group(1).strip()
    return None


def talep_bilgisi(metin: str | None) -> dict:
    if not metin:
        return {"unvan": None, "vkn": None, "limit": None}
    limit = None
    t = alan(metin, "Talep Edilen Toplam Limit", "Toplam Limit", "Talep Edilen Limit")
    if t:
        m = re.search(r"\d[\d.,]*", t)
        limit = para(m.group(0)) if m else None
    if limit is None:
        t = alan(metin, "Talep") or ""
        limit = sum((para(x) or SIFIR for x in re.findall(r"(\d{1,3}(?:\.\d{3})+(?:,\d+)?)\s*TL", t)), SIFIR) or None
    vkn = alan(metin, "VKN", "Vergi No", "Vergi Kimlik No")
    return {"unvan": alan(metin, "Firma Unvanı", "Unvan", "Firma"), "vkn": re.sub(r"\D", "", vkn) if vkn else None, "limit": limit}


# ----------------------------------------------------------------------------
# Kod hesapları ve kontroller
# ----------------------------------------------------------------------------

def _bol(a, b):
    return a / b if a is not None and b not in (None, 0) else None


def rasyolar(yillar: list[str], veri: dict) -> dict[str, dict]:
    sonuc = {}
    onceki = None
    for y in yillar:
        v = veri[y]
        g = lambda k: v.get(k)  # noqa: E731
        favok = g("favok")
        if favok is None and g("brut") is not None and g("fgid") is not None:
            favok = g("brut") - g("fgid") + (g("amort") or SIFIR)
        finborc = (g("kvbk") or SIFIR) + (g("uvbk") or SIFIR) if (g("kvbk") is not None or g("uvbk") is not None) else None
        netborc = finborc - (g("hazir") or SIFIR) if finborc is not None else None
        r = {"favok": favok, "finborc": finborc, "netborc": netborc,
             "cari": _bol(g("dv"), g("kvyk")),
             "asit": _bol(g("dv") - (g("stok") or SIFIR), g("kvyk")) if g("dv") is not None else None,
             "kaldirac": _bol((g("kvyk") or SIFIR) + (g("uvyk") or SIFIR), g("aktif")),
             "ozk_aktif": _bol(g("ozk"), g("aktif")),
             "nb_favok": _bol(netborc, favok) if favok and favok > 0 else None,
             "faiz_karsilama": _bol(favok, g("fin")),
             "favok_marji": _bol(favok, g("satis")), "net_marj": _bol(g("net"), g("satis")),
             "alacak_gun": _bol(g("alacak") * 365, g("satis")) if g("alacak") is not None else None,
             "stok_gun": _bol(g("stok") * 365, g("smm")) if g("stok") is not None else None,
             "buyume": _bol(g("satis") - onceki["satis"], onceki["satis"]) if onceki and g("satis") is not None and onceki.get("satis") else None,
             "fin_artis": _bol(g("fin") - onceki["fin"], onceki["fin"]) if onceki and g("fin") is not None and onceki.get("fin") else None,
             "net_degisim": _bol(g("net") - onceki["net"], abs(onceki["net"])) if onceki and g("net") is not None and onceki.get("net") else None,
             "denklik": (g("aktif") - (g("kvyk") or SIFIR) - (g("uvyk") or SIFIR) - g("ozk"))
             if g("aktif") is not None and g("ozk") is not None else None}
        sonuc[y] = r
        onceki = v
    return sonuc


def kontrol_et(girdi: dict, metinler: list[Dosya], katsayilar: dict) -> dict:
    bulgular = []

    def b(onem, aciklama, kaynak=""):
        bulgular.append((onem, aciklama, kaynak))

    talep = talep_bilgisi(girdi["talep"])
    if girdi["talep"] is None:
        b("yüksek", "Talep ve firma bilgisi dosyası (adında 'talep' geçen .txt/.docx/.pdf) yok")
    elif talep["limit"] is None:
        b("orta", "Talep edilen toplam limit okunamadı ('Talep Edilen Toplam Limit: … TL' satırı ekleyin)")
    r, yillar = {}, []
    if girdi["mali"]:
        yillar, veri = girdi["mali"]
        r = rasyolar(yillar, veri)
        son = r[yillar[-1]]
        y = yillar[-1]
        v = veri[y]
        if son["denklik"] and abs(son["denklik"]) > 1:
            b("yüksek", f"{y} bilançosu denk değil: aktif − (KVYK + UVYK + özkaynak) = {tl(son['denklik'])} TL", "mali tablolar")
        if v.get("ozk") is not None and v["ozk"] <= 0:
            b("yüksek", f"{y} özkaynak negatif veya sıfır ({tl(v['ozk'])} TL): teknik iflas / TTK 376 değerlendirmesi", "mali tablolar")
        if son["cari"] is not None and son["cari"] < 1:
            b("yüksek", f"{y} cari oran {oran(son['cari'])} (< 1): kısa vadeli borçlar dönen varlıkları aşıyor", "mali tablolar")
        elif son["cari"] is not None and son["cari"] < Decimal("1.2"):
            b("orta", f"{y} cari oran {oran(son['cari'])} (< 1,20)", "mali tablolar")
        if son["asit"] is not None and son["asit"] < 1:
            b("bilgi", f"{y} asit-test oranı {oran(son['asit'])}: likidite stoklara bağlı", "mali tablolar")
        if son["kaldirac"] is not None and son["kaldirac"] > Decimal("0.7"):
            b("yüksek", f"{y} kaldıraç {yuzde(son['kaldirac'])} (> %70)", "mali tablolar")
        if son["favok"] is not None and son["favok"] <= 0:
            b("yüksek", f"{y} FAVÖK negatif ({tl(son['favok'])} TL): borç servisi faaliyetten karşılanamıyor", "mali tablolar")
        elif son["nb_favok"] is not None and son["nb_favok"] > 4:
            b("yüksek", f"{y} net finansal borç / FAVÖK {oran(son['nb_favok'])} (> 4)", "mali tablolar")
        elif son["nb_favok"] is not None and son["nb_favok"] > 3:
            b("orta", f"{y} net finansal borç / FAVÖK {oran(son['nb_favok'])} (> 3)", "mali tablolar")
        if son["faiz_karsilama"] is not None and son["faiz_karsilama"] < Decimal("1.5"):
            b("yüksek", f"{y} faiz karşılama (FAVÖK / finansman gideri) {oran(son['faiz_karsilama'])} (< 1,5)", "mali tablolar")
        if v.get("net") is not None and v["net"] < 0:
            b("yüksek", f"{y} dönem net zararı {tl(v['net'])} TL", "mali tablolar")
        if son["buyume"] is not None and son["buyume"] < 0:
            b("orta", f"{y} net satışlar {yuzde(-son['buyume'])} azaldı", "mali tablolar")
        if son["fin_artis"] is not None and son["buyume"] is not None and son["fin_artis"] > son["buyume"] + Decimal("0.25"):
            b("orta", f"{y} finansman giderleri {yuzde(son['fin_artis'])} arttı, net satışlar {yuzde(son['buyume'])}: borçlanma maliyeti "
                      "cirodan hızlı büyüyor", "mali tablolar")
        if son["net_degisim"] is not None and son["net_degisim"] < Decimal("-0.2"):
            b("orta", f"{y} net kâr {yuzde(-son['net_degisim'])} düştü", "mali tablolar")
        if len(yillar) > 1:
            o = r[yillar[-2]]
            if son["alacak_gun"] is not None and o["alacak_gun"] is not None and son["alacak_gun"] - o["alacak_gun"] > 15:
                b("orta", f"alacak tahsil süresi {o['alacak_gun']:.0f} günden {son['alacak_gun']:.0f} güne uzadı", "mali tablolar")
            if son["stok_gun"] is not None and o["stok_gun"] is not None and son["stok_gun"] - o["stok_gun"] > 15:
                b("orta", f"stok devir süresi {o['stok_gun']:.0f} günden {son['stok_gun']:.0f} güne uzadı", "mali tablolar")
    else:
        b("yüksek", "Mali tablo dosyası yok (Kalem × yıl tablosu: Dönen Varlıklar, Toplam Aktif, Net Satışlar…)")

    kkb = girdi["kkb"] or []
    risk = {"nl": sum((x["nl"] for x in kkb), SIFIR), "nr": sum((x["nr"] for x in kkb), SIFIR),
            "gl": sum((x["gl"] for x in kkb), SIFIR), "gr": sum((x["gr"] for x in kkb), SIFIR)}
    if kkb:
        dol = _bol(risk["nr"], risk["nl"])
        if dol is not None and dol > Decimal("0.9"):
            b("orta", f"sektörde nakdi limit doluluğu {yuzde(dol)} (risk {tl(risk['nr'])} / limit {tl(risk['nl'])} TL)", "KKB")
        for x in kkb:
            if x["nr"] > x["nl"] > 0 or x["gr"] > x["gl"] > 0:
                b("yüksek", f"{x['banka']}: risk limiti aşıyor", "KKB")
        if r:
            fb = r[yillar[-1]]["finborc"]
            if fb and abs(risk["nr"] - fb) / fb > Decimal("0.2"):
                b("orta", f"KKB nakdi risk ({tl(risk['nr'])} TL) ile {yillar[-1]} bilançosundaki banka kredileri ({tl(fb)} TL) arasında "
                          f"{yuzde(abs(risk['nr'] - fb) / fb)} fark: dönem farkı mı, bilançoda gösterilmeyen borç mu?", "KKB, mali tablolar")
            satis = girdi["mali"][1][yillar[-1]].get("satis")
            if satis and talep["limit"] and (risk["nr"] + talep["limit"]) / satis > Decimal("0.5"):
                b("orta", f"mevcut nakdi risk + talep, son yıl net satışların {yuzde((risk['nr'] + talep['limit']) / satis)}'i", "KKB, talep")
    else:
        b("orta", "KKB / memzuç risk dosyası yok: sektördeki risk ve limitler görülmeden teklif hazırlanıyor")

    tem = girdi["teminat"] or []
    toplam_tem = SIFIR
    for t in tem:
        k = katla(t["tur"])
        t["katsayi"] = next((v for ad, v in sorted(katsayilar.items(), key=lambda x: -len(x[0])) if ad in k), None)
        t["kabul"] = (t["deger"] or SIFIR) * t["katsayi"] if t["katsayi"] is not None else None
        if t["katsayi"] is None:
            b("orta", f"teminat türü '{t['tur']}' katsayı listesinde yok; karşılamaya dahil edilmedi", "teminatlar")
        toplam_tem += t["kabul"] or SIFIR
        if "ipotek" in k and not re.search(r"ekspertiz|değerleme", t["aciklama"], re.I):
            b("orta", f"ipotek değeri için ekspertiz/değerleme tarihi belirtilmemiş: {t['aciklama'] or t['tur']}", "teminatlar")
    karsilama = _bol(toplam_tem, talep["limit"])
    if tem and karsilama is not None and karsilama < 1:
        b("orta" if karsilama >= Decimal("0.5") else "yüksek",
          f"teminat karşılama oranı {yuzde(karsilama)} (katsayılı teminat {tl(toplam_tem)} TL / talep {tl(talep['limit'])} TL)", "teminatlar")
    if not tem:
        b("orta", "Teminat dosyası yok")
    elif not toplam_tem:
        b("yüksek", "Katsayılı teminat değeri sıfır (yalnız şahsi kefalet vb.): teklif fiilen teminatsız", "teminatlar")

    for d in metinler:
        if d.tur not in ("İstihbarat", "Diğer not"):
            continue
        for cumle in re.split(r"(?<=[.!?])\s+|\n", d.metin):
            ck = kucuk(cumle)
            for kelime, onem in OLUMSUZ:
                if kelime in ck and not re.search(r"\b(yok|bulunmamaktadır|bulunmamakta|rastlanmamıştır|görülmemiştir)\b", ck):
                    b(onem, f"istihbaratta '{kelime}': {cumle.strip()[:160]}", d.ad)
                    break
    return {"bulgular": bulgular, "talep": talep, "rasyolar": r, "yillar": yillar, "kkb": kkb, "risk": risk, "teminat": tem,
            "teminat_toplam": toplam_tem, "karsilama": karsilama}


# ----------------------------------------------------------------------------
# Maskeleme
# ----------------------------------------------------------------------------

def takma_adlar(talep: dict, terimler: list[str]) -> dict[str, str]:
    harita = {}
    if talep.get("unvan"):
        harita[talep["unvan"]] = "[FİRMA]"
        kok = re.split(r"\s+(?:San\.|Sanayi|Tic\.|Ticaret|A\.Ş\.|Ltd\.)", talep["unvan"])[0].strip()
        if len(kok) > 4 and kok != talep["unvan"]:
            harita[kok] = "[FİRMA]"
    if talep.get("vkn"):
        harita[talep["vkn"]] = "[VKN]"
    for i, t in enumerate(sorted({t.strip() for t in terimler if t.strip()}, key=len, reverse=True), 1):
        harita[t] = f"[GİZLİ-{i}]"
        soyad = t.split()[-1]
        if len(t.split()) > 1 and len(soyad) > 3 and soyad not in harita:
            harita[soyad] = f"[GİZLİ-{i}-SOYAD]"
    return harita


def maskele(metin: str, harita: dict[str, str]) -> str:
    for gercek, takma in sorted(harita.items(), key=lambda x: -len(x[0])):
        metin = re.sub(re.escape(gercek), takma, metin, flags=re.I)
    return llm.maskele(metin)


def geri_ac(metin: str, harita: dict[str, str]) -> str:
    for gercek, takma in sorted(harita.items(), key=lambda x: -len(x[1])):
        if takma == "[FİRMA]":
            continue
        metin = metin.replace(takma, gercek)
    unvan = next((g for g, t in harita.items() if t == "[FİRMA]"), None)
    return metin.replace("[FİRMA]", unvan) if unvan else metin


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

SEMA = {
    "type": "object",
    "properties": {
        "firma_ve_talep": {"type": "string"},
        "mali_degerlendirme": {"type": "string"},
        "istihbarat_degerlendirmesi": {"type": "string"},
        "teminat_degerlendirmesi": {"type": "string"},
        "guclu_yonler": {"type": "array", "items": {"type": "string"}},
        "riskler": {"type": "array", "items": {
            "type": "object", "properties": {"aciklama": {"type": "string"}, "kaynak": {"type": "string"}},
            "required": ["aciklama", "kaynak"], "additionalProperties": False}},
        "onerilen_sartlar": {"type": "array", "items": {"type": "string"}},
        "eksik_bilgiler": {"type": "array", "items": {"type": "string"}},
        "teklif_ozeti": {"type": "string"},
    },
    "required": ["firma_ve_talep", "mali_degerlendirme", "istihbarat_degerlendirmesi", "teminat_degerlendirmesi", "guclu_yonler",
                 "riskler", "onerilen_sartlar", "eksik_bilgiler", "teklif_ozeti"],
    "additionalProperties": False,
}
METIN_ALANLARI = ("firma_ve_talep", "mali_degerlendirme", "istihbarat_degerlendirmesi", "teminat_degerlendirmesi", "teklif_ozeti")


def kod_ozeti(k: dict) -> list[str]:
    s = [f"- Talep edilen toplam limit: {tl(k['talep']['limit'])} TL"]
    for y in k["yillar"]:
        r = k["rasyolar"][y]
        s.append(f"- {y}: cari oran {oran(r['cari'])}, asit-test {oran(r['asit'])}, kaldıraç {yuzde(r['kaldirac'])}, "
                 f"özkaynak/aktif {yuzde(r['ozk_aktif'])}, FAVÖK {tl(r['favok'])} TL, FAVÖK marjı {yuzde(r['favok_marji'])}, "
                 f"net kâr marjı {yuzde(r['net_marj'])}, finansal borç {tl(r['finborc'])} TL, net finansal borç {tl(r['netborc'])} TL, "
                 f"net finansal borç/FAVÖK {oran(r['nb_favok'])}, faiz karşılama {oran(r['faiz_karsilama'])}, "
                 f"alacak tahsil süresi {oran(r['alacak_gun'], 0)} gün, stok devir süresi {oran(r['stok_gun'], 0)} gün, "
                 f"ciro büyümesi {yuzde(r['buyume'])}")
    if k["kkb"]:
        x = k["risk"]
        s.append(f"- KKB toplam: nakdi limit {tl(x['nl'])} / risk {tl(x['nr'])} TL, gayrinakdi limit {tl(x['gl'])} / risk {tl(x['gr'])} TL, "
                 f"{len(k['kkb'])} kurum")
    if k["teminat"]:
        for t in k["teminat"]:
            s.append(f"- Teminat: {t['tur']} — {t['aciklama']} — değer {tl(t['deger'])} TL × katsayı {oran(t['katsayi'])} = {tl(t['kabul'])} TL")
        s.append(f"- Katsayılı teminat toplamı {tl(k['teminat_toplam'])} TL, karşılama oranı {yuzde(k['karsilama'])}")
    return s


def sayilar(metin: str) -> set[str]:
    """3+ haneli sayıların yalnız rakamları ('15.000.000' → '15000000', '1,25' → '125')."""
    return {re.sub(r"\D", "", m) for m in re.findall(r"\d[\d.,]*\d", metin) if len(re.sub(r"\D", "", m)) >= 3}


def hazirla(metinler: list[Dosya], k: dict, harita: dict) -> tuple[str, str]:
    kk = [f"- [{o}] {maskele(a, harita)} (kaynak: {kay or '-'})" for o, a, kay in k["bulgular"]]
    mesaj = "\n".join([
        "<belgeler>",
        *[f'<belge ad="{d.ad}" tur="{d.tur}">\n{maskele(d.metin, harita)}\n</belge>' for d in metinler if d.metin],
        "</belgeler>",
        "<kod_hesaplari>", *[maskele(x, harita) for x in kod_ozeti(k)], "</kod_hesaplari>",
        "<kod_kontrolleri>", *(kk or ["- Kontrol bulgusu yok"]), "</kod_kontrolleri>",
    ])
    return (BURASI / "prompt.md").read_text(encoding="utf-8"), mesaj


def model_yaz(metinler, k, harita) -> dict:
    sistem, mesaj = hazirla(metinler, k, harita)
    yanit = llm.json_iste(sistem, mesaj, SEMA)
    bilinen = sayilar(mesaj)
    supheli = []
    for alan_ad in METIN_ALANLARI:
        supheli += [s for s in sayilar(yanit[alan_ad]) if s not in bilinen]
    for liste in ("guclu_yonler", "onerilen_sartlar", "eksik_bilgiler"):
        for x in yanit[liste]:
            supheli += [s for s in sayilar(x) if s not in bilinen]
    for x in yanit["riskler"]:
        supheli += [s for s in sayilar(x["aciklama"]) if s not in bilinen]
    yanit["dogrulanamayan_sayilar"] = sorted(set(supheli))
    return yanit


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
ONEM_DOLGU = {"yüksek": PatternFill("solid", fgColor="FDE2E1"), "orta": PatternFill("solid", fgColor="FFF4CE"),
              "bilgi": PatternFill("solid", fgColor="E8F0FE"), "model": PatternFill("solid", fgColor="F2F2F2")}
ONAY_DOLGU = PatternFill("solid", fgColor="FFF4CE")
UST = Alignment(vertical="top", wrap_text=True)
RASYO_SATIR = [("Cari oran", "cari", oran), ("Asit-test oranı", "asit", oran), ("Kaldıraç (yabancı kaynak / aktif)", "kaldirac", yuzde),
               ("Özkaynak / aktif", "ozk_aktif", yuzde), ("FAVÖK", "favok", tl), ("FAVÖK marjı", "favok_marji", yuzde),
               ("Net kâr marjı", "net_marj", yuzde), ("Finansal borç", "finborc", tl), ("Net finansal borç", "netborc", tl),
               ("Net finansal borç / FAVÖK", "nb_favok", oran), ("Faiz karşılama", "faiz_karsilama", oran),
               ("Alacak tahsil süresi (gün)", "alacak_gun", lambda x: oran(x, 0)), ("Stok devir süresi (gün)", "stok_gun", lambda x: oran(x, 0)),
               ("Net satış büyümesi", "buyume", yuzde)]


def _baslik(ws):
    for h in ws[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI


def rapor_yaz(cikti: Path, klasor: Path, metinler, atlanan, k: dict, y: dict, harita: dict) -> Path:
    g = lambda s: geri_ac(s, harita)  # noqa: E731
    unvan = k["talep"]["unvan"] or klasor.name
    md = [f"# Kredi teklif özeti — {unvan}", "",
          f"**Talep edilen toplam limit:** {tl(k['talep']['limit'])} TL · **Teminat karşılama:** {yuzde(k['karsilama'])}"
          f" · **Hazırlanma:** {datetime.now():%d.%m.%Y}", "",
          "## Firma ve talep", g(y["firma_ve_talep"]), "", "## Mali değerlendirme", g(y["mali_degerlendirme"]), ""]
    if k["yillar"]:
        md += ["| Rasyo | " + " | ".join(k["yillar"]) + " |", "|---|" + "---:|" * len(k["yillar"])]
        md += [f"| {ad} | " + " | ".join(f(k["rasyolar"][yy][kd]) for yy in k["yillar"]) + " |" for ad, kd, f in RASYO_SATIR]
        md.append("")
    if k["kkb"]:
        md += ["**KKB / memzuç**", "", "| Kurum | Nakdi limit | Nakdi risk | Gayrinakdi limit | Gayrinakdi risk |", "|---|---:|---:|---:|---:|"]
        md += [f"| {x['banka']} | {tl(x['nl'])} | {tl(x['nr'])} | {tl(x['gl'])} | {tl(x['gr'])} |" for x in k["kkb"]]
        r = k["risk"]
        md += [f"| **Toplam** | **{tl(r['nl'])}** | **{tl(r['nr'])}** | **{tl(r['gl'])}** | **{tl(r['gr'])}** |", ""]
    md += ["## İstihbarat", g(y["istihbarat_degerlendirmesi"]), "", "## Teminatlar", g(y["teminat_degerlendirmesi"]), ""]
    if k["teminat"]:
        md += ["| Tür | Açıklama | Değer | Katsayı | Kabul edilen |", "|---|---|---:|---:|---:|"]
        md += [f"| {t['tur']} | {g(t['aciklama'])} | {tl(t['deger'])} | {oran(t['katsayi'])} | {tl(t['kabul'])} |" for t in k["teminat"]]
        md += [f"| **Toplam** | | | | **{tl(k['teminat_toplam'])}** |", ""]
    md += ["## Güçlü yönler", *[f"- {g(x)}" for x in y["guclu_yonler"]], "", "## Riskler ve kontrol bulguları"]
    md += [f"- **[{o}]** {g(a)}" + (f" _({kay})_" if kay else "") + " — _kod kontrolü_" for o, a, kay in k["bulgular"]]
    md += [f"- {g(x['aciklama'])} _({x['kaynak']})_" for x in y["riskler"]]
    md += ["", "## Önerilen şartlar (taslak)", *[f"- {g(x)}" for x in y["onerilen_sartlar"]], "",
           "## Eksik bilgi ve belgeler", *[f"- {g(x)}" for x in y["eksik_bilgiler"]], "",
           "## Portföy yöneticisi görüşü (taslak)", g(y["teklif_ozeti"]), ""]
    if y["dogrulanamayan_sayilar"]:
        md += [f"> **Dikkat:** Model metnindeki şu sayılar girdilerde ve kod hesaplarında bulunamadı, kontrol edin: "
               f"{', '.join(y['dogrulanamayan_sayilar'])}", ""]
    md += ["---", f"_Kaynaklar: {', '.join(f'{d.ad} ({d.tur})' for d in metinler)}_" + (f"  \n_Okunamayan: {', '.join(atlanan)}_" if atlanan else ""),
           "_Yapay zekâ destekli taslak. Rasyolar, risk ve teminat tabloları koddan; değerlendirmeler taslaktır. Teklif portföy "
           "yöneticisinin kontrolü ve imzasıyla tahsis birimine gönderilir; kredi kararı yetkili organlarındır._"]
    cikti.parent.mkdir(parents=True, exist_ok=True)
    md_yol = cikti.with_suffix(".md")
    md_yol.write_text("\n".join(md), encoding="utf-8")

    wb = Workbook()
    s = wb.active
    s.title = "Özet"
    for satir in [["Kredi teklif özeti", unvan], ["Talep edilen limit", tl(k["talep"]["limit"]) + " TL"],
                  ["Teminat karşılama", yuzde(k["karsilama"])], [], ["Firma ve talep", g(y["firma_ve_talep"])],
                  ["Mali değerlendirme", g(y["mali_degerlendirme"])], ["İstihbarat", g(y["istihbarat_degerlendirmesi"])],
                  ["Teminatlar", g(y["teminat_degerlendirmesi"])], ["Güçlü yönler", "\n".join(f"• {g(x)}" for x in y["guclu_yonler"])],
                  ["Önerilen şartlar", "\n".join(f"• {g(x)}" for x in y["onerilen_sartlar"])],
                  ["Eksik bilgiler", "\n".join(f"• {g(x)}" for x in y["eksik_bilgiler"])], ["Görüş (taslak)", g(y["teklif_ozeti"])],
                  ["Doğrulanamayan sayılar", ", ".join(y["dogrulanamayan_sayilar"]) or "—"],
                  ["Model", llm.kullanim_ozeti()], ["Oluşturulma", datetime.now().strftime("%d.%m.%Y %H:%M")]]:
        s.append(satir)
    s["A1"].font = Font(bold=True, size=13)
    s.column_dimensions["A"].width = 24
    s.column_dimensions["B"].width = 120
    for r in s.iter_rows():
        for h in r:
            h.alignment = UST

    ws = wb.create_sheet("Kontroller")
    ws.append(["Önem", "Bulgu", "Kaynak", "Üreten", "Portföy Yöneticisi Onayı"])
    _baslik(ws)
    for o, a, kay in k["bulgular"]:
        ws.append([o, g(a), kay, "Kod", ""])
        ws.cell(ws.max_row, 1).fill = ONEM_DOLGU[o]
        ws.cell(ws.max_row, 5).fill = ONAY_DOLGU
    for x in y["riskler"]:
        ws.append(["model", g(x["aciklama"]), x["kaynak"], "Model", ""])
        ws.cell(ws.max_row, 1).fill = ONEM_DOLGU["model"]
        ws.cell(ws.max_row, 5).fill = ONAY_DOLGU
    for sy in y["dogrulanamayan_sayilar"]:
        ws.append(["orta", f"Model metninde girdilerde olmayan sayı: {sy}", "model metni", "Kod", ""])
        ws.cell(ws.max_row, 1).fill = ONEM_DOLGU["orta"]
        ws.cell(ws.max_row, 5).fill = ONAY_DOLGU
    for j, w in enumerate((9, 100, 22, 8, 16), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for r in ws.iter_rows():
        for h in r:
            h.alignment = UST

    if k["yillar"]:
        ra = wb.create_sheet("Rasyolar")
        ra.append(["Rasyo"] + k["yillar"])
        _baslik(ra)
        for ad, kd, _ in RASYO_SATIR:
            ra.append([ad] + [None if k["rasyolar"][yy][kd] is None else float(k["rasyolar"][yy][kd]) for yy in k["yillar"]])
            fmt = "0.0%" if kd in ("kaldirac", "ozk_aktif", "favok_marji", "net_marj", "buyume") else ("#,##0" if kd in ("favok", "finborc", "netborc") else "0.00")
            for c in ra[ra.max_row][1:]:
                c.number_format = fmt
        ra.column_dimensions["A"].width = 34
    if k["kkb"]:
        kk = wb.create_sheet("KKB Risk")
        kk.append(["Kurum", "Nakdi Limit", "Nakdi Risk", "Gayrinakdi Limit", "Gayrinakdi Risk", "Nakdi Doluluk"])
        _baslik(kk)
        for x in k["kkb"] + [{"banka": "TOPLAM", **k["risk"]}]:
            kk.append([x["banka"], float(x["nl"]), float(x["nr"]), float(x["gl"]), float(x["gr"]),
                       float(x["nr"] / x["nl"]) if x["nl"] else None])
            for c in kk[kk.max_row][1:5]:
                c.number_format = "#,##0"
            kk.cell(kk.max_row, 6).number_format = "0.0%"
        for j, w in enumerate((18, 15, 15, 16, 16, 12), 1):
            kk.column_dimensions[get_column_letter(j)].width = w
    if k["teminat"]:
        te = wb.create_sheet("Teminatlar")
        te.append(["Tür", "Açıklama", "Değer", "Katsayı", "Kabul Edilen"])
        _baslik(te)
        for t in k["teminat"]:
            te.append([t["tur"], g(t["aciklama"]), float(t["deger"]) if t["deger"] is not None else None,
                       float(t["katsayi"]) if t["katsayi"] is not None else None, float(t["kabul"]) if t["kabul"] is not None else None])
            te.cell(te.max_row, 3).number_format = te.cell(te.max_row, 5).number_format = "#,##0"
        te.append(["TOPLAM", "", None, None, float(k["teminat_toplam"])])
        te.cell(te.max_row, 5).number_format = "#,##0"
        te.append(["Not", "Katsayılar örnektir; bankanızın teminat politikasındaki oranlarla --katsayilar dosyasından değiştirin."])
        for j, w in enumerate((18, 60, 15, 9, 15), 1):
            te.column_dimensions[get_column_letter(j)].width = w
    wb.save(cikti)
    return md_yol


def katsayi_oku(yol: Path | None) -> dict:
    if not yol:
        return dict(VARSAYILAN_KATSAYI)
    s = tablo_oku(yol)
    k = {}
    for r in s[1:]:
        if r and r[0] and len(r) > 1 and para(r[1]) is not None:
            v = para(r[1])
            k[katla(r[0])] = v / 100 if v > 1 else v
    return k


def calistir(klasor: Path, cikti: Path, terimler: list[str] | None = None, katsayi_yolu: Path | None = None, evet: bool = False) -> dict:
    if not klasor.is_dir():
        raise llm.LLMHatasi(f"{klasor}: klasör bulunamadı (teklif belgelerini bir klasörde toplayın).")
    girdi, metinler, atlanan = dosyalari_oku(klasor)
    if not metinler:
        raise llm.LLMHatasi(f"{klasor}: okunabilir dosya yok.")
    k = kontrol_et(girdi, metinler, katsayi_oku(katsayi_yolu))
    print(f"[OK] {len(metinler)} dosya · kod kontrolü: {len(k['bulgular'])} bulgu · teminat karşılama {yuzde(k['karsilama'])}")
    for o, a, _ in k["bulgular"]:
        print(f"[!] [{o}] {a}")
    harita = takma_adlar(k["talep"], terimler or [])
    llm.onay_al(f"{len(metinler)} dosyanın metni ve kod hesapları (firma unvanı, VKN ve {len(terimler or [])} gizli terim takma adlı; "
                "telefon/e-posta/TCKN/IBAN maskeli) teklif özeti için gönderilecek. Bankacılık sırrı ve iç politikanızı kontrol edin.", evet)
    y = model_yaz(metinler, k, harita)
    md = rapor_yaz(cikti, klasor, metinler, atlanan, k, y, harita)
    return {"girdi": girdi, "metinler": metinler, "kontrol": k, "yanit": y, "md": md, "harita": harita}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    ornek = BURASI / "ornek_veri" / "teklif_ornek_ambalaj"
    p = argparse.ArgumentParser(description="Firma bilgisi, mali tablolar, KKB, istihbarat ve teminatlardan kredi teklif özeti hazırlar.")
    p.add_argument("--girdi", type=Path, default=ornek, help="Teklif klasörü (talep .txt, mali tablolar, KKB, teminatlar, istihbarat)")
    p.add_argument("--gizle", nargs="*", default=[], metavar="TERİM", help="Modele gönderilmeden maskelenecek kişi adları (ortaklar, kefiller)")
    p.add_argument("--katsayilar", type=Path, help="Teminat türü → katsayı (.csv: Tür;Katsayı). Varsayılanlar örnektir.")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kredi_teklif_ozeti.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    if a.girdi == ornek:
        a.gizle = a.gizle or ["Kemal Örnekoğlu", "Selin Örnekoğlu"]
    try:
        s = calistir(a.girdi, a.cikti, a.gizle, a.katsayilar, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    if s["yanit"]["dogrulanamayan_sayilar"]:
        print(f"[!] Model metninde girdilerde olmayan sayılar: {', '.join(s['yanit']['dogrulanamayan_sayilar'])}")
    print(f"[OK] Teklif özeti: {s['md'].resolve()}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
