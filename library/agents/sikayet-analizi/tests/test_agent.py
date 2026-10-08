"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import re
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri" / "sikayetler.csv"
T, A, H, I, O, F, C, S, P = ("Teslimat gecikmesi", "Ürün arızası", "Hasarlı / eksik / yanlış ürün", "İade ve para iadesi", "Ödeme ve fatura",
                             "Fiyat ve kampanya", "Müşteri hizmetlerine ulaşım", "Teknik servis", "Personel davranışı")
KONULAR = [T, A, H, I, O, F, C, S, P]
ATAMA = {  # no → (konu, kök neden)
    1: (A, "Ürün kalitesi / kusur"), 2: (T, "Teslimat / lojistik"), 3: (O, "Sistem / teknik arıza"), 4: (C, "Bilgilendirme / iletişim"),
    5: (I, "Bilgilendirme / iletişim"), 6: (S, "Tedarikçi / iş ortağı"), 7: (A, "Ürün kalitesi / kusur"), 8: (H, "Süreç / prosedür"),
    9: (F, "Sistem / teknik arıza"), 10: (P, "Personel davranışı"), 11: (T, "Tedarikçi / iş ortağı"), 12: (A, "Ürün kalitesi / kusur"),
    13: (A, "Ürün kalitesi / kusur"), 14: (T, "Teslimat / lojistik"), 15: (I, "Süreç / prosedür"), 16: (O, "Fiyat / ücret / fatura"),
    17: (H, "Teslimat / lojistik"), 18: (C, "Bilgilendirme / iletişim"), 19: (S, "Tedarikçi / iş ortağı"), 20: (A, "Ürün kalitesi / kusur"),
    21: (P, "Personel davranışı"), 22: (T, "Tedarikçi / iş ortağı"), 23: (O, "Süreç / prosedür"), 24: (T, "Teslimat / lojistik"),
    25: (T, "Teslimat / lojistik"), 26: (T, "Bilgilendirme / iletişim"), 27: (A, "Ürün kalitesi / kusur"), 28: (T, "Teslimat / lojistik"),
    29: (I, "Süreç / prosedür"), 30: (H, "Teslimat / lojistik"), 31: (F, "Sistem / teknik arıza"), 32: (T, "Tedarikçi / iş ortağı"),
    33: (A, "Ürün kalitesi / kusur"), 34: (H, "Süreç / prosedür"), 35: (T, "Teslimat / lojistik"), 36: (T, "Sistem / teknik arıza"),
    37: (C, "Bilgilendirme / iletişim"), 38: ("Kargo şubesi yoğunluğu", "Tedarikçi / iş ortağı"),   # listede yok → Diğer
}
RISK = {6: ["Hukuki süreç (hakem heyeti, dava, avukat)"], 12: ["Sağlık / güvenlik"], 23: ["Kişisel veri"],
        26: ["Sosyal medya / basın"], 35: ["Hukuki süreç (hakem heyeti, dava, avukat)"], 10: ["Kaba davranış / ayrımcılık", "Uydurma risk"]}
TEKRAR = {7, 20, 33}
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append((sistem, kullanici))
    if "konu_onerisi" in kullanici:
        return {"konular": [{"ad": k, "tanim": f"{k} şikâyetleri"} for k in KONULAR] + [{"ad": "teslimat  gecikmesi", "tanim": "tekrar"},
                                                                                     {"ad": "Diğer", "tanim": "x"}]}
    nolar = re.findall(r'<kayit no="([^"]+)"', kullanici)
    yanit = []
    for no in nolar:
        n = int(no.split("-")[1])
        yanit.append({"no": no, "konu": ATAMA[n][0], "urun": "", "kok_neden": ATAMA[n][1], "kok_neden_aciklama": "Metinde yazıyor.",
                      "ozet": f"özet {no}", "riskler": RISK.get(n, []), "tekrar_belirtiyor": n in TEKRAR})
    return {"kayitlar": yanit + [{"no": "S-999", "konu": T, "urun": "", "kok_neden": "Belirsiz", "kok_neden_aciklama": "", "ozet": "",
                                  "riskler": [], "tekrar_belirtiyor": False}]}


class KodTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.k, cls.u = agent.kayitlari_oku(ORNEK)
        agent.kod_hesapla(cls.k, date(2026, 9, 28), 15)
        cls.d = {x["no"]: x for x in cls.k}

    def test_okuma(self):
        self.assertEqual(len(self.k), 38)
        self.assertEqual(self.u, [])
        self.assertEqual(sum(1 for x in self.k if not x["kapali"]), 8)
        self.assertEqual(self.d["S-006"]["sure"], 21)
        self.assertEqual(self.d["S-022"]["kanal"], "Web Formu")
        self.assertIn('"teslim edildi"', self.d["S-022"]["metin"])

    def test_hedef_ve_tekrar(self):
        self.assertEqual(self.d["S-006"]["notlar"], ["21 günde kapandı (hedef 15 gün)"])
        self.assertEqual(self.d["S-029"]["notlar"], ["Açık, 16 gündür bekliyor (hedef 15 gün)"])
        self.assertEqual(self.d["S-007"]["notlar"], ["Aynı müşterinin 15 gün önceki şikâyeti: S-001"])
        self.assertEqual(self.d["S-034"]["notlar"], ["Aynı müşterinin 21 gün önceki şikâyeti: S-024"])
        self.assertIn("Sonraki 30 gün içinde tekrar şikâyet: S-034", self.d["S-024"]["notlar"])
        self.assertEqual(self.d["S-033"]["notlar"], [])                  # farklı müşteri, metindeki "yine" modelden gelir

    def test_maskeleme(self):
        m = agent.maskeli_metin(self.d["S-001"], [])
        self.assertTrue(m.startswith("Ben [MÜŞTERİ], aldığım"))
        m = agent.maskeli_metin(self.d["S-026"], ["sosyal medya"])
        self.assertIn("[TELEFON]", m)
        self.assertIn("[GİZLİ-1]", m)

    def test_egilim(self):
        self.assertEqual(agent.egilim_durumu([2, 2, 8], 3), (2, "Yükselen"))
        self.assertEqual(agent.egilim_durumu([0, 0, 3], 3), (0, "Yeni"))
        self.assertEqual(agent.egilim_durumu([6, 6, 2], 3), (6, "Azalan"))
        self.assertEqual(agent.egilim_durumu([2, 2, 2], 3)[1], "Sabit")
        self.assertEqual(agent.egilim_durumu([5], 3), (None, "—"))
        self.assertEqual(len(agent.donemler(date(2025, 11, 20), date(2026, 2, 1), "ay")), 4)
        self.assertEqual(len(agent.donemler(date(2026, 9, 1), date(2026, 9, 28), "hafta")), 5)

    def test_baslik_satiri_ve_konu_dosyasi(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp) / "s.csv"
            t.write_text("Şikâyet Raporu 2026\n\nKAYIT NO,AÇIKLAMA,DURUM\n1,Kargo gelmedi,Çözüldü\n2,Ürün bozuk,\n", encoding="utf-8")
            k, _ = agent.kayitlari_oku(t)
            (Path(tmp) / "k.txt").write_text("# konular\n- Kargo: kargo sorunları\nArıza\nDiğer: x\n", encoding="utf-8")
            konular = agent.konular_oku(Path(tmp) / "k.txt")
        self.assertEqual([(x["no"], x["kapali"]) for x in k], [("1", True), ("2", False)])
        self.assertEqual([(x["ad"], x["tanim"]) for x in konular], [("Kargo", "kargo sorunları"), ("Arıza", "")])


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_uctan_uca(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "s.xlsx"
            s = agent.calistir(ORNEK, cikti, evet=True)
            self.assertEqual(len(MESAJLAR), 3)                            # konu önerisi + 2 paket (25 + 13)
            self.assertIn("GÖREV 1", MESAJLAR[0][0])
            self.assertNotIn("GÖREV 2", MESAJLAR[0][0])
            self.assertIn("GÖREV 2", MESAJLAR[1][0])
            self.assertNotIn("GÖREV 1", MESAJLAR[1][0])
            tum = "".join(m for _, m in MESAJLAR)
            for gizli in ("Elif", "Kurgu", "0532", "M1001"):
                self.assertNotIn(gizli, tum)
            self.assertEqual([k["ad"] for k in s["konular"]], KONULAR)     # yinelenen ve "Diğer" önerisi atıldı
            self.assertNotIn("S-999", s["sonuc"])
            self.assertEqual(s["sonuc"]["S-038"]["konu"], "Diğer")
            self.assertEqual(s["sonuc"]["S-010"]["riskler"], ["Kaba davranış / ayrımcılık"])
            self.assertTrue(any("S-038" in u for u in s["uyarilar"]))
            tablo = {x["konu"]: x for x in s["ozet"]["tablo"]}
            tes = tablo[T]
            self.assertEqual((tes["adet"], tes["seri"], tes["egilim"], tes["acik"]), (11, [2, 2, 7], "Yükselen", 3))
            self.assertEqual(s["ozet"]["tablo"][0]["konu"], T)
            self.assertEqual(tablo[A]["seri"], [3, 2, 2])
            self.assertEqual(tablo[A]["egilim"], "Sabit")
            self.assertTrue(any("eksik olabilir" in u for u in s["ozet"]["uyarilar"]))
            self.assertEqual(s["ozet"]["urun"]["Kahve Makinesi K200"][A], 3)
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Özet", "Konu Analizi", "Eğilim", "Kök Neden", "Takip Listesi", "Kayıtlar", "Konu Listesi"])
            ozet = [[c.value for c in r] for r in wb["Özet"].iter_rows()]
            self.assertIn(["Yükselen", f"{T}: son dönem 7 kayıt, önceki dönem ortalaması 2.0"], ozet)
            takip = {r[0].value for r in wb["Takip Listesi"].iter_rows(min_row=2)}
            self.assertEqual(takip, {"S-006", "S-010", "S-012", "S-023", "S-026", "S-029", "S-033", "S-034", "S-035"})
            ws = wb["Kayıtlar"]
            self.assertEqual(ws.cell(1, 16).value, "Kontrol / Düzeltme")
            self.assertEqual(ws.max_row, 39)

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_konu_dosyasi_ile_oneri_yok(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            (Path(tmp) / "k.txt").write_text("\n".join(f"{k}: tanım" for k in KONULAR), encoding="utf-8")
            s = agent.calistir(ORNEK, Path(tmp) / "s.xlsx", Path(tmp) / "k.txt", tur="hafta", evet=True)
        self.assertEqual(len(MESAJLAR), 2)
        self.assertEqual(s["konular"][0]["kaynak"], "Kullanıcı")
        self.assertEqual(len(s["ozet"]["donemler"]), 14)

    def test_dosya_yok(self):
        with self.assertRaises(llm.LLMHatasi):
            agent.calistir(Path("olmayan.csv"), Path("x.xlsx"), evet=True)


if __name__ == "__main__":
    unittest.main()
