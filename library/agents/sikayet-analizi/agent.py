"""
Şikâyet Analizi — Workers / Workless AI Agent
Müşteri Hizmetleri ve Çağrı Merkezi › Müşteri Deneyimi Uzmanı

1. Kod şikâyet kayıtlarını okur; çözüm süresini, açık kayıtların yaşını (hedef süre aşımı), aynı müşterinin
   30 gün içinde tekrarlayan şikâyetlerini ve kanal dağılımını hesaplar. Müşteri adları metinde [MÜŞTERİ]
   olur; telefon, e-posta, TCKN ve IBAN maskelenir.
2. Konu listesi (--konular) verilmezse model, kayıtlardan bir örneklemle 5-12 konu başlığı önerir (kümeleme).
3. Model her şikâyeti bu konulardan birine atar; ürünü, olası kök neden kategorisini ve risk işaretlerini
   (hakem heyeti/dava, resmî kurum, sosyal medya, sağlık-güvenlik, kişisel veri, kaba davranış) çıkarır.
4. Kod eğilimi hesaplar: konu × dönem tablosu, son dönemin önceki üç dönem ortalamasına göre değişimi
   ("yükselen" konular), Pareto (80/20), konu × kök neden ve ürün × konu tabloları.

Kullanım:
    python agent.py                                              # örnek: 38 şikâyet, Temmuz-Eylül 2026
    python agent.py --girdi sikayetler.xlsx --konular konular.txt --donem hafta --hedef-gun 10
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import llm

BURASI = Path(__file__).resolve().parent
PAKET = 25
ORNEKLEM = 150
TEKRAR_GUN = 30
DIGER = "Diğer"
KOK_NEDENLER = ["Ürün kalitesi / kusur", "Teslimat / lojistik", "Fiyat / ücret / fatura", "Personel davranışı", "Süreç / prosedür",
                "Bilgilendirme / iletişim", "Sistem / teknik arıza", "Tedarikçi / iş ortağı", "Müşteri beklentisi / kullanım", "Belirsiz"]
RISKLER = ["Hukuki süreç (hakem heyeti, dava, avukat)", "Resmî kurum / regülatör başvurusu", "Sosyal medya / basın", "Sağlık / güvenlik",
           "Kişisel veri", "Kaba davranış / ayrımcılık"]
SUTUNLAR = {
    "no": ("sikayet no", "kayit no", "no", "id", "talep no", "basvuru no", "sikayet id"),
    "tarih": ("tarih", "acilis tarihi", "kayit tarihi", "sikayet tarihi", "olusturma tarihi", "basvuru tarihi"),
    "kanal": ("kanal", "kaynak", "basvuru kanali"),
    "musteri": ("musteri no", "musteri kodu", "musteri id", "musteri"),
    "musteri_ad": ("musteri adi", "ad soyad", "musteri adi soyadi", "adi soyadi"),
    "urun": ("urun", "urun/hizmet", "urun / hizmet", "hizmet", "urun grubu", "urun adi"),
    "metin": ("sikayet", "sikayet metni", "aciklama", "metin", "mesaj", "icerik", "sikayet detayi"),
    "durum": ("durum", "statu", "kayit durumu"),
    "kapanis": ("kapanis tarihi", "cozum tarihi", "kapanma tarihi", "kapatilma tarihi"),
}
KAPALI = {"kapandi", "kapali", "cozuldu", "tamamlandi", "sonuclandi", "kapatildi"}
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")


def katla(s) -> str:
    return " ".join(str(s or "").translate(_TR).lower().split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


# ----------------------------------------------------------------------------
# Kayıtlar ve kod hesapları
# ----------------------------------------------------------------------------

def kayitlari_oku(yol: Path) -> tuple[list[dict], list[str]]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:10]):                       # başlık satırını bul (üstteki rapor başlıklarını atla)
        b = [katla(x) for x in r]
        k = {ad: next((i for i, x in enumerate(b) if x in es), None) for ad, es in SUTUNLAR.items()}
        if k["metin"] is not None:
            break
    else:
        raise llm.LLMHatasi(f"{yol.name}: 'Şikâyet' (metin) sütunu bulunamadı. İlk satır: {s[0] if s else '(boş)'}")
    al = lambda r, a: r[k[a]] if k[a] is not None and k[a] < len(r) else None  # noqa: E731
    kayitlar, uyarilar = [], []
    for i, r in enumerate(s[bi + 1:], 1):
        metin = str(al(r, "metin") or "").strip()
        if not metin:
            continue
        no = str(al(r, "no") or f"K{i}").strip()
        t = tarih(al(r, "tarih"))
        if t is None:
            uyarilar.append(f"{no}: tarih okunamadı ({al(r, 'tarih')!r}); eğilim ve süre hesaplarına girmez")
        kap = tarih(al(r, "kapanis"))
        durum = str(al(r, "durum") or "").strip()
        kapali = kap is not None or katla(durum) in KAPALI
        kayitlar.append({"no": no, "tarih": t, "kanal": str(al(r, "kanal") or "").strip() or "(belirtilmemiş)",
                         "musteri": str(al(r, "musteri") or "").strip(), "musteri_ad": str(al(r, "musteri_ad") or "").strip(),
                         "urun": str(al(r, "urun") or "").strip(), "metin": metin, "durum": durum or ("Kapandı" if kapali else "Açık"),
                         "kapali": kapali, "kapanis": kap,
                         "sure": (kap - t).days if kap and t else None})
        if kap and t and kap < t:
            uyarilar.append(f"{no}: kapanış tarihi açılıştan önce")
    sayac = Counter(k_["no"] for k_ in kayitlar)
    uyarilar += [f"{n}: kayıt numarası {c} kez geçiyor" for n, c in sayac.items() if c > 1]
    return kayitlar, uyarilar


def kod_hesapla(kayitlar: list[dict], rapor_tarihi: date, hedef_gun: int) -> None:
    """Hedef süre aşımı ve 30 gün içinde tekrar eden müşteri şikâyetlerini kayıtlara işler."""
    for k in kayitlar:
        k["notlar"] = []
        if not k["kapali"] and k["tarih"]:
            k["yas"] = (rapor_tarihi - k["tarih"]).days
            if k["yas"] > hedef_gun:
                k["notlar"].append(f"Açık, {k['yas']} gündür bekliyor (hedef {hedef_gun} gün)")
        elif k["sure"] is not None and k["sure"] > hedef_gun:
            k["notlar"].append(f"{k['sure']} günde kapandı (hedef {hedef_gun} gün)")
    musteri = defaultdict(list)
    for k in kayitlar:
        anahtar = k["musteri"] or k["musteri_ad"]
        if anahtar and k["tarih"]:
            musteri[katla(anahtar)].append(k)
    for ks in musteri.values():
        ks.sort(key=lambda x: x["tarih"])
        for a, b in zip(ks, ks[1:]):
            if (b["tarih"] - a["tarih"]).days <= TEKRAR_GUN:
                b["notlar"].append(f"Aynı müşterinin {(b['tarih'] - a['tarih']).days} gün önceki şikâyeti: {a['no']}")
                if not any(n.startswith("Sonraki") for n in a["notlar"]):
                    a["notlar"].append(f"Sonraki {TEKRAR_GUN} gün içinde tekrar şikâyet: {b['no']}")


def maskeli_metin(k: dict, gizli: list[str]) -> str:
    metin = k["metin"]
    adlar = [k["musteri_ad"]] + [p for p in k["musteri_ad"].split() if len(p) >= 3] if k["musteri_ad"] else []
    for ad in sorted(set(adlar), key=len, reverse=True):
        metin = re.sub(rf"(?<!\w){re.escape(ad)}(?!\w)", "[MÜŞTERİ]", metin, flags=re.I)
    metin = re.sub(r"\[MÜŞTERİ\](\s+\[MÜŞTERİ\])+", "[MÜŞTERİ]", metin)
    for i, g in enumerate(gizli, 1):
        metin = re.sub(re.escape(g), f"[GİZLİ-{i}]", metin, flags=re.I)
    return llm.maskele(metin)


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

TAKSONOMI_SEMA = {
    "type": "object",
    "properties": {"konular": {"type": "array", "items": {
        "type": "object", "properties": {"ad": {"type": "string"}, "tanim": {"type": "string"}},
        "required": ["ad", "tanim"], "additionalProperties": False}}},
    "required": ["konular"], "additionalProperties": False,
}


def siniflama_semasi(konular: list[str]) -> dict:
    return {
        "type": "object",
        "properties": {"kayitlar": {"type": "array", "items": {
            "type": "object",
            "properties": {
                "no": {"type": "string"},
                "konu": {"type": "string", "enum": konular},
                "urun": {"type": "string"},
                "kok_neden": {"type": "string", "enum": KOK_NEDENLER},
                "kok_neden_aciklama": {"type": "string"},
                "ozet": {"type": "string"},
                "riskler": {"type": "array", "items": {"type": "string", "enum": RISKLER}},
                "tekrar_belirtiyor": {"type": "boolean"},
            },
            "required": ["no", "konu", "urun", "kok_neden", "kok_neden_aciklama", "ozet", "riskler", "tekrar_belirtiyor"],
            "additionalProperties": False}}},
        "required": ["kayitlar"], "additionalProperties": False,
    }


def _kayit_xml(k: dict, gizli: list[str]) -> str:
    urun = f' urun="{k["urun"]}"' if k["urun"] else ""
    return f'<kayit no="{k["no"]}" kanal="{k["kanal"]}"{urun}>\n{maskeli_metin(k, gizli)}\n</kayit>'


def konu_oner(kayitlar: list[dict], gizli: list[str]) -> list[dict]:
    """Konu listesi verilmediyse modelden örneklemle konu başlıkları (kümeler) ister."""
    adim = max(1, len(kayitlar) // ORNEKLEM)
    ornek = kayitlar[::adim][:ORNEKLEM]
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8").split("## GÖREV 2")[0]
    mesaj = "\n".join(["<gorev>konu_onerisi</gorev>", f"<kayitlar adet=\"{len(ornek)}\" toplam=\"{len(kayitlar)}\">",
                       *[_kayit_xml(k, gizli) for k in ornek], "</kayitlar>"])
    yanit = llm.json_iste(sistem, mesaj, TAKSONOMI_SEMA)
    konular, gorulen = [], set()
    for x in yanit.get("konular", []):
        ad = " ".join(str(x.get("ad", "")).split())
        if ad and katla(ad) not in gorulen and katla(ad) != katla(DIGER):
            gorulen.add(katla(ad))
            konular.append({"ad": ad, "tanim": str(x.get("tanim", "")).strip(), "kaynak": "Model önerisi"})
    if not konular:
        raise llm.LLMHatasi("Model konu listesi önermedi; --konular ile bir liste verin.")
    return konular[:12]


def siniflandir(kayitlar: list[dict], konular: list[dict], gizli: list[str]) -> tuple[dict[str, dict], list[str]]:
    adlar = [k["ad"] for k in konular] + [DIGER]
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    sistem = sistem.split("## GÖREV 1")[0] + "## GÖREV 2" + sistem.split("## GÖREV 2")[1]
    liste = "\n".join(f"- {k['ad']}: {k['tanim']}" for k in konular) + f"\n- {DIGER}: listedeki hiçbir konuya uymayanlar"
    sonuc, uyarilar = {}, []
    sema = siniflama_semasi(adlar)
    for i in range(0, len(kayitlar), PAKET):
        parca = kayitlar[i:i + PAKET]
        mesaj = "\n".join([f"<konu_listesi>\n{liste}\n</konu_listesi>", "<kayitlar>", *[_kayit_xml(k, gizli) for k in parca], "</kayitlar>"])
        yanit = llm.json_iste(sistem, mesaj, sema)
        gecerli = {k["no"] for k in parca}
        for s in yanit.get("kayitlar", []):
            if s.get("no") not in gecerli:
                continue
            if s.get("konu") not in adlar:
                uyarilar.append(f"{s['no']}: model listede olmayan konu verdi ({s.get('konu')!r}); '{DIGER}' sayıldı")
                s["konu"] = DIGER
            if s.get("kok_neden") not in KOK_NEDENLER:
                s["kok_neden"] = "Belirsiz"
            s["riskler"] = [r for r in s.get("riskler", []) if r in RISKLER]
            sonuc[s["no"]] = s
    eksik = [k["no"] for k in kayitlar if k["no"] not in sonuc]
    if eksik:
        uyarilar.append(f"{len(eksik)} kayıt için model yanıt vermedi: {', '.join(eksik[:10])}{'…' if len(eksik) > 10 else ''}")
    return sonuc, uyarilar


def konular_oku(yol: Path) -> list[dict]:
    konular = []
    for satir in yol.read_text(encoding="utf-8-sig").splitlines():
        ad, _, tanim = satir.partition(":")
        ad = " ".join(ad.strip(" -•\t").split())
        if ad and not ad.startswith("#") and katla(ad) != katla(DIGER):
            konular.append({"ad": ad, "tanim": tanim.strip(), "kaynak": "Kullanıcı"})
    if not konular:
        raise llm.LLMHatasi(f"{yol.name}: konu bulunamadı (satır başına 'Konu: tanım').")
    return konular


# ----------------------------------------------------------------------------
# Eğilim ve özet (kod)
# ----------------------------------------------------------------------------

def donem_anahtari(t: date, tur: str):
    return (t.year, t.month) if tur == "ay" else t - timedelta(days=t.weekday())


def donem_adi(d, tur: str) -> str:
    return f"{d[1]:02d}.{d[0]}" if tur == "ay" else f"{d:%d.%m.%Y} haftası"


def donemler(ilk: date, son: date, tur: str) -> list:
    d, sonuc = donem_anahtari(ilk, tur), []
    bitis = donem_anahtari(son, tur)
    while d <= bitis:
        sonuc.append(d)
        d = ((d[0] + d[1] // 12, d[1] % 12 + 1) if tur == "ay" else d + timedelta(days=7))
    return sonuc


def egilim_durumu(seri: list[int], esik: int) -> tuple[float | None, str]:
    son, onceki = seri[-1], seri[-4:-1]
    if not onceki:
        return None, "—"
    ort = sum(onceki) / len(onceki)
    if son >= esik and son >= 1.5 * ort:
        return ort, "Yükselen" if ort else "Yeni"
    if ort >= esik and son <= 0.5 * ort:
        return ort, "Azalan"
    return ort, "Sabit"


def ozetle(kayitlar: list[dict], sonuc: dict, konular: list[dict], tur: str, rapor_tarihi: date, esik: int) -> dict:
    adlar = [k["ad"] for k in konular] + [DIGER]
    tarihli = [k for k in kayitlar if k["tarih"] and k["no"] in sonuc]
    dl = donemler(min(k["tarih"] for k in tarihli), max(k["tarih"] for k in tarihli), tur) if tarihli else []
    matris = {a: Counter() for a in adlar}
    for k in tarihli:
        matris[sonuc[k["no"]]["konu"]][donem_anahtari(k["tarih"], tur)] += 1
    say = Counter(sonuc[k["no"]]["konu"] for k in kayitlar if k["no"] in sonuc)
    toplam = sum(say.values())
    tablo, kum = [], 0
    for ad, n in sorted(say.items(), key=lambda x: (-x[1], adlar.index(x[0]))):
        kum += n
        ks = [k for k in kayitlar if k["no"] in sonuc and sonuc[k["no"]]["konu"] == ad]
        sureler = [k["sure"] for k in ks if k["sure"] is not None]
        seri = [matris[ad][d] for d in dl]
        ort, durum = egilim_durumu(seri, esik) if seri else (None, "—")
        tablo.append({"konu": ad, "adet": n, "pay": n / toplam, "kumulatif": kum / toplam, "pareto": "A" if (kum - n) / toplam < 0.8 else "",
                      "acik": sum(1 for k in ks if not k["kapali"]), "ort_sure": sum(sureler) / len(sureler) if sureler else None,
                      "hedef_asimi": sum(1 for k in ks if k["notlar"] and "hedef" in " ".join(k["notlar"])),
                      "seri": seri, "son": seri[-1] if seri else 0, "onceki_ort": ort, "egilim": durum,
                      "kok": Counter(sonuc[k["no"]]["kok_neden"] for k in ks)})
    uyarilar = []
    if dl:
        son = max(k["tarih"] for k in tarihli)
        bitis = (date(son.year + son.month // 12, son.month % 12 + 1, 1) - timedelta(days=1)) if tur == "ay" else dl[-1] + timedelta(days=6)
        if son < bitis:
            uyarilar.append(f"Son dönem ({donem_adi(dl[-1], tur)}) eksik olabilir: kayıtlar {son:%d.%m.%Y} tarihine kadar. "
                            "Bu dönemdeki artış gerçekte daha da yüksek olabilir.")
        if len(dl) < 2:
            uyarilar.append("Eğilim için en az iki dönem gerekir.")
    urun = defaultdict(Counter)
    for k in kayitlar:
        if k["no"] in sonuc:
            u = k["urun"] or sonuc[k["no"]]["urun"].strip() or "(belirtilmemiş)"
            urun[u][sonuc[k["no"]]["konu"]] += 1
    return {"tablo": tablo, "donemler": dl, "urun": urun, "uyarilar": uyarilar,
            "kanal": Counter(k["kanal"] for k in kayitlar), "kok": Counter(s["kok_neden"] for s in sonuc.values()),
            "risk": Counter(r for s in sonuc.values() for r in s["riskler"])}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
ONAY_DOLGU = PatternFill("solid", fgColor="FFF4CE")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
KALIN = Font(bold=True)


def _baslik(ws, basliklar, genislik=None):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for j, w in enumerate(genislik or (), 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def rapor_yaz(cikti: Path, kayitlar: list[dict], sonuc: dict, konular: list[dict], ozet: dict, tur: str, rapor_tarihi: date,
              hedef_gun: int, uyarilar: list[str]) -> None:
    wb = Workbook()
    o = wb.active
    o.title = "Özet"
    o.column_dimensions["A"].width, o.column_dimensions["B"].width = 34, 90
    acik = [k for k in kayitlar if not k["kapali"]]
    sureler = [k["sure"] for k in kayitlar if k["sure"] is not None]
    for etiket, deger in [("Rapor tarihi", f"{rapor_tarihi:%d.%m.%Y}"), ("Toplam şikâyet", len(kayitlar)),
                          ("Açık", len(acik)), (f"Hedef süreyi ({hedef_gun} gün) aşan",
                                                sum(1 for k in kayitlar if any("hedef" in n for n in k["notlar"]))),
                          ("Ortalama çözüm süresi (gün)", round(sum(sureler) / len(sureler), 1) if sureler else "-"),
                          ("Tekrarlayan (aynı müşteri, 30 gün)", sum(1 for k in kayitlar if any(n.startswith("Aynı müşterinin") for n in k["notlar"]))),
                          ("Kanallar", ", ".join(f"{a} {n}" for a, n in ozet["kanal"].most_common())),
                          ("Olası kök nedenler (model)", ", ".join(f"{a} {n}" for a, n in ozet["kok"].most_common())),
                          ("Risk işaretleri (model)", ", ".join(f"{a} {n}" for a, n in ozet["risk"].most_common()) or "-")]:
        o.append([etiket, deger])
        o.cell(o.max_row, 1).font = KALIN
    o.append([])
    _baslik(o, ["Eğilim", "Konu ve gerekçe"])
    for x in ozet["tablo"]:
        if x["egilim"] in ("Yükselen", "Yeni", "Azalan"):
            o.append([x["egilim"], f"{x['konu']}: son dönem {x['son']} kayıt, önceki dönem ortalaması "
                                   f"{0 if x['onceki_ort'] is None else round(x['onceki_ort'], 1)}"])
            o.cell(o.max_row, 1).fill = YESIL if x["egilim"] == "Azalan" else KIRMIZI
    o.append([])
    for u in ozet["uyarilar"] + uyarilar:
        o.append(["Uyarı", u])
        o.cell(o.max_row, 1).fill = SARI
    o.append(["Model", llm.kullanim_ozeti()])
    o.append(["Not", "Sayımlar, süreler, eğilim ve tekrar tespiti koddan; konu, ürün, kök neden ve risk işaretleri modelden gelir "
                     "(mor hücreler). Kök neden metinden çıkarılan bir öneridir; kesin kök neden analizi (5 Neden, balık kılçığı) "
                     "ilgili ekiple yapılmalıdır."])
    for r in o.iter_rows():
        for h in r:
            h.alignment = UST

    k = wb.create_sheet("Konu Analizi")
    _baslik(k, ["Konu", "Adet", "Pay", "Kümülatif", "Pareto", "Açık", "Ort. Çözüm (gün)", "Hedef Aşımı", "Son Dönem",
                "Önceki Ort.", "Eğilim", "Başlıca Kök Neden (model)"], (30, 8, 8, 10, 8, 7, 14, 11, 10, 11, 11, 40))
    for x in ozet["tablo"]:
        k.append([x["konu"], x["adet"], x["pay"], x["kumulatif"], x["pareto"], x["acik"],
                  None if x["ort_sure"] is None else round(x["ort_sure"], 1), x["hedef_asimi"], x["son"],
                  None if x["onceki_ort"] is None else round(x["onceki_ort"], 1), x["egilim"],
                  ", ".join(f"{a} ({n})" for a, n in x["kok"].most_common(2))])
        for c in (3, 4):
            k.cell(k.max_row, c).number_format = "0.0%"
        if x["egilim"] in ("Yükselen", "Yeni"):
            k.cell(k.max_row, 11).fill = KIRMIZI
        k.cell(k.max_row, 12).fill = AI_DOLGU
    n = len(ozet["tablo"])
    if n:
        g = BarChart()
        g.title, g.height, g.width = "Konu bazında şikâyet (Pareto)", 9, 18
        g.add_data(Reference(k, min_col=2, min_row=1, max_row=1 + n), titles_from_data=True)
        g.set_categories(Reference(k, min_col=1, min_row=2, max_row=1 + n))
        k.add_chart(g, "N2")
    k.freeze_panes = "B2"

    e = wb.create_sheet("Eğilim")
    dl = ozet["donemler"]
    _baslik(e, ["Konu"] + [donem_adi(d, tur) for d in dl] + ["Eğilim"], (30,) + (16,) * len(dl) + (11,))
    for x in ozet["tablo"]:
        e.append([x["konu"], *x["seri"], x["egilim"]])
        if x["egilim"] in ("Yükselen", "Yeni"):
            e.cell(e.max_row, len(dl) + 2).fill = KIRMIZI
    e.append(["Toplam", *[sum(x["seri"][i] for x in ozet["tablo"]) for i in range(len(dl))], ""])
    for h in e[e.max_row]:
        h.font = KALIN
    if dl and n:
        c = LineChart()
        c.title, c.height, c.width = "İlk 5 konu — dönem bazında", 9, 18
        for i in range(min(5, n)):
            c.add_data(Reference(e, min_col=1, max_col=1 + len(dl), min_row=2 + i), from_rows=True, titles_from_data=True)
        c.set_categories(Reference(e, min_col=2, max_col=1 + len(dl), min_row=1))
        e.add_chart(c, f"{get_column_letter(len(dl) + 4)}2")

    kn = wb.create_sheet("Kök Neden")
    kokler = [x for x in KOK_NEDENLER if ozet["kok"][x]]
    _baslik(kn, ["Konu \\ Olası kök neden (model)"] + kokler, (30,) + (16,) * len(kokler))
    for x in ozet["tablo"]:
        kn.append([x["konu"]] + [x["kok"][a] or None for a in kokler])
    kn.append([])
    urun_konular = [x["konu"] for x in ozet["tablo"]]
    _baslik(kn, ["Ürün / hizmet \\ Konu"] + urun_konular + ["Toplam"])
    for u, c_ in sorted(ozet["urun"].items(), key=lambda i: -sum(i[1].values())):
        kn.append([u] + [c_[a] or None for a in urun_konular] + [sum(c_.values())])

    r = wb.create_sheet("Takip Listesi")
    _baslik(r, ["Şikâyet No", "Tarih", "Durum", "Konu", "Risk İşaretleri (model)", "Kod Kontrolleri", "Özet (model)", "Yapılacak / Sorumlu"],
            (12, 12, 10, 26, 34, 46, 50, 30))
    for kayit in sorted(kayitlar, key=lambda x: (x["kapali"], x["tarih"] or date.min)):
        s = sonuc.get(kayit["no"], {})
        if not (s.get("riskler") or kayit["notlar"] or s.get("tekrar_belirtiyor")) or (kayit["kapali"] and not s.get("riskler")):
            continue
        notlar = kayit["notlar"] + (["Müşteri daha önce de şikâyet ettiğini belirtiyor"] if s.get("tekrar_belirtiyor") else [])
        r.append([kayit["no"], kayit["tarih"], kayit["durum"], s.get("konu", ""), "\n".join(s.get("riskler", [])), "\n".join(notlar),
                  s.get("ozet", ""), ""])
        r.cell(r.max_row, 2).number_format = "DD.MM.YYYY"
        if s.get("riskler"):
            r.cell(r.max_row, 5).fill = KIRMIZI
        for c in (4, 5, 7):
            if c != 5 or not s.get("riskler"):
                r.cell(r.max_row, c).fill = AI_DOLGU
        r.cell(r.max_row, 8).fill = ONAY_DOLGU
        for h in r[r.max_row]:
            h.alignment = UST

    d = wb.create_sheet("Kayıtlar")
    _baslik(d, ["Şikâyet No", "Tarih", "Kanal", "Müşteri", "Ürün", "Şikâyet", "Durum", "Kapanış", "Çözüm (gün)", "Konu (model)",
                "Ürün (model)", "Olası Kök Neden (model)", "Kök Neden Açıklaması (model)", "Risk (model)", "Kod Kontrolleri",
                "Kontrol / Düzeltme"], (11, 11, 13, 10, 20, 60, 9, 11, 9, 24, 18, 22, 40, 26, 36, 18))
    for kayit in kayitlar:
        s = sonuc.get(kayit["no"])
        d.append([kayit["no"], kayit["tarih"], kayit["kanal"], kayit["musteri"] or kayit["musteri_ad"], kayit["urun"], kayit["metin"],
                  kayit["durum"], kayit["kapanis"], kayit["sure"],
                  *(([s["konu"], s["urun"], s["kok_neden"], s["kok_neden_aciklama"], "\n".join(s["riskler"])]) if s else
                    ["AI yanıt vermedi", "", "", "", ""]), "\n".join(kayit["notlar"]), ""])
        for c in (2, 8):
            d.cell(d.max_row, c).number_format = "DD.MM.YYYY"
        for c in range(10, 15):
            d.cell(d.max_row, c).fill = AI_DOLGU
        d.cell(d.max_row, 16).fill = ONAY_DOLGU
        if kayit["notlar"]:
            d.cell(d.max_row, 15).fill = SARI
        for h in d[d.max_row]:
            h.alignment = UST
    d.freeze_panes = "B2"
    d.auto_filter.ref = d.dimensions

    t = wb.create_sheet("Konu Listesi")
    _baslik(t, ["Konu", "Tanım", "Kaynak"], (30, 80, 16))
    for x in konular:
        t.append([x["ad"], x["tanim"], x["kaynak"]])
        if x["kaynak"] == "Model önerisi":
            t.cell(t.max_row, 1).fill = AI_DOLGU
    t.append([DIGER, "Listedeki hiçbir konuya uymayanlar", "Sabit"])
    t.append([])
    t.append(["İpucu", "Listeyi düzeltip satır başına 'Konu: tanım' olarak bir .txt dosyasına yazın ve --konular ile verin; "
                       "böylece dönemler arasında aynı konu başlıkları kullanılır ve eğilim karşılaştırılabilir olur."])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


# ----------------------------------------------------------------------------
# Akış
# ----------------------------------------------------------------------------

def calistir(girdi: Path, cikti: Path, konu_dosyasi: Path | None = None, tur: str = "ay", hedef_gun: int = 15, esik: int = 3,
             rapor_tarihi: date | None = None, gizli: list[str] | None = None, evet: bool = False) -> dict:
    if not girdi.exists():
        raise llm.LLMHatasi(f"{girdi} bulunamadı.")
    kayitlar, uyarilar = kayitlari_oku(girdi)
    if not kayitlar:
        raise llm.LLMHatasi(f"{girdi.name}: şikâyet kaydı bulunamadı.")
    tarihler = [x for k in kayitlar for x in (k["tarih"], k["kapanis"]) if x]
    rapor_tarihi = rapor_tarihi or (max(tarihler) if tarihler else date.today())
    kod_hesapla(kayitlar, rapor_tarihi, hedef_gun)
    konular = konular_oku(konu_dosyasi) if konu_dosyasi else None
    print(f"[OK] {len(kayitlar)} şikâyet · açık {sum(1 for k in kayitlar if not k['kapali'])} · rapor tarihi {rapor_tarihi:%d.%m.%Y}")
    istek = "konu önerisi + sınıflandırma" if konular is None else "sınıflandırma"
    llm.onay_al(f"{len(kayitlar)} şikâyet metni ({istek}) gönderilecek. Müşteri adları [MÜŞTERİ], {len(gizli or [])} gizli terim "
                "takma adlı; telefon, e-posta, TCKN ve IBAN maskeli. Müşteri numarası gönderilmez.", evet)
    if konular is None:
        konular = konu_oner(kayitlar, gizli or [])
        print(f"[OK] Model {len(konular)} konu önerdi: " + "; ".join(k["ad"] for k in konular))
    sonuc, model_uyarilari = siniflandir(kayitlar, konular, gizli or [])
    ozet = ozetle(kayitlar, sonuc, konular, tur, rapor_tarihi, esik)
    rapor_yaz(cikti, kayitlar, sonuc, konular, ozet, tur, rapor_tarihi, hedef_gun, uyarilar + model_uyarilari)
    return {"kayitlar": kayitlar, "sonuc": sonuc, "konular": konular, "ozet": ozet, "uyarilar": uyarilar + model_uyarilari}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Şikâyet kayıtlarını konu, ürün ve kök nedene göre kümeleyip eğilim raporu çıkarır.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "sikayetler.csv",
                   help="Şikâyet kayıtları (.xlsx/.csv): Şikâyet (metin) zorunlu; No, Tarih, Kanal, Müşteri No/Adı, Ürün, Durum, Kapanış Tarihi")
    p.add_argument("--konular", type=Path, help="Sabit konu listesi (.txt, satır başına 'Konu: tanım'); verilmezse model önerir")
    p.add_argument("--donem", choices=["ay", "hafta"], default="ay", help="Eğilim dönemi (varsayılan: ay)")
    p.add_argument("--hedef-gun", type=int, default=15, help="Şikâyet çözüm hedef süresi, takvim günü (varsayılan 15, örnektir)")
    p.add_argument("--esik", type=int, default=3, help="'Yükselen' sayılması için son dönemdeki en az kayıt (varsayılan 3)")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY (varsayılan: verideki son tarih)")
    p.add_argument("--gizle", nargs="*", default=[], metavar="TERİM", help="Modele gönderilmeden takma adla değiştirilecek adlar")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "sikayet_analizi.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    rt = None
    if a.tarih:
        rt = tarih(a.tarih)
        if rt is None:
            print(f"[X] Tarih GG.AA.YYYY olmalı: {a.tarih}")
            return 2
    try:
        s = calistir(a.girdi, a.cikti, a.konular, a.donem, a.hedef_gun, a.esik, rt, a.gizle, a.evet)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(s['sonuc'])}/{len(s['kayitlar'])} kayıt sınıflandırıldı")
    for x in s["ozet"]["tablo"][:5]:
        print(f"     {x['konu']}: {x['adet']} (%{x['pay'] * 100:.0f}) · eğilim: {x['egilim']}")
    for u in s["ozet"]["uyarilar"] + s["uyarilar"]:
        print(f"[!] {u}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
