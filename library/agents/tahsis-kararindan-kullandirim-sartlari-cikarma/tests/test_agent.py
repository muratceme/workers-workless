"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import json
import os
import re
import sys
import tempfile
import unittest
from datetime import datetime
from decimal import Decimal as D
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri" / "tahsis_kararlari.txt"
CEVAP = json.loads((Path(__file__).parent / "sahte_cevaplar.json").read_text(encoding="utf-8"))
ADLAR = ["Hasan Demirtaş", "Selin Demirtaş"]
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    """ON-6101'de ilk turda 4. maddeyi (müşteri çeki temliki) atlar."""
    MESAJLAR.append(kullanici)
    no = re.search(r'<karar no="([^"]+)"', kullanici)[1]
    return copy.deepcopy(CEVAP[f"{no}_hatali" if no == "ON-6101" and "<duzeltme>" not in kullanici else no])


def kararlar():
    return {k["no"]: k for k in agent.kararlari_oku(ORNEK, agent.takma_adlar(ADLAR))}


class KodTesti(unittest.TestCase):
    def test_okuma(self):
        k = kararlar()
        self.assertEqual(list(k), ["ON-6101", "ON-6102"])
        self.assertIn("[KİŞİ-1] ve [KİŞİ-2]'ın", k["ON-6101"]["metin"])
        self.assertNotIn("Demirtaş", k["ON-6101"]["metin"])
        self.assertTrue(k["ON-6102"]["metin"].rstrip().endswith("yazılar alınacaktır."))
        self.assertEqual((agent.para("25.000.000"), agent.para("3.000"), agent.para("1.250,50"), agent.para("41")), (D(25000000), D(3000), D("1250.50"), D(41)))
        with tempfile.TemporaryDirectory() as tmp:
            (Path(tmp) / "a.txt").write_text("Karar No: X-1\nLimit 1.000.000 TL\n", encoding="utf-8")
            (Path(tmp) / "b.txt").write_text("Karar No: X-1\nLimit 2.000.000 TL\n", encoding="utf-8")
            with self.assertRaises(llm.LLMHatasi):
                agent.kararlari_oku(Path(tmp), {})
            (Path(tmp) / "b.txt").unlink()
            (Path(tmp) / "c.txt").write_text("Limit 2.000.000 TL\n", encoding="utf-8")
            self.assertEqual([x["no"] for x in agent.kararlari_oku(Path(tmp), {})], ["X-1", "c"])   # başlıksız dosya: dosya adı

    def test_denetim(self):
        k = kararlar()
        self.assertEqual([agent.denetle(k["ON-6101"], CEVAP["ON-6101"]), agent.denetle(k["ON-6102"], CEVAP["ON-6102"])], [[], []])
        self.assertEqual(agent.denetle(k["ON-6101"], CEVAP["ON-6101_hatali"]),
                         ["ON-6101: şart içeren cümle şartlara alınmamış: \"Rotatif kredi kullandırımlarında limitin %20'si oranında müşteri çeki temlik alınacaktır.\""])
        c = copy.deepcopy(CEVAP["ON-6102"])
        c["sartlar"][0]["gereken_tutar"] = "5.000.000"                   # %80 → 4.800.000 olmalı
        c["sartlar"][1]["zamanlama"] = "Kullandırım sonrası"              # cümlede sonraya işaret yok
        c["sartlar"][2]["kaynak"] = "SGK ve vergi borcu yoktur yazısı alınacaktır."  # kararda aynen yok
        c["azami_vade_ay"] = "36"
        s = agent.denetle(k["ON-6102"], c)
        for parca in ("kararda olmayan sayı: 36, 5000000", "yüzdeli tutarı boş bırakın", "'Kullandırım öncesi' seçin", "kaynak cümle kararda aynen geçmiyor",
                      "şart içeren cümle şartlara alınmamış: \"- Kullandırımdan önce"):
            self.assertTrue(any(parca in x for x in s), parca)
        c = copy.deepcopy(CEVAP["ON-6101"])
        c["sartlar"][2]["zamanlama"] = "Kullandırım öncesi"
        c["sartlar"][4]["termin"] = ""
        c["sartlar"][0]["gereken_tutar"] = ""
        c["urunler"].append("Spot")
        s = agent.denetle(k["ON-6101"], c)
        for parca in ("'takiben'", "termin boş", "gereken tutar boş", "ürün 'Spot'"):
            self.assertTrue(any(parca in x for x in s), parca)

    def test_tamamla(self):
        c = agent.tamamla(copy.deepcopy(CEVAP["ON-6102"]))
        self.assertEqual((c["sartlar"][0]["gereken_tutar"], c["sartlar"][0]["hesap"]), ("4.800.000,00", "Kod hesapladı: limit 6.000.000,00 × %80"))
        self.assertEqual(c["kod_uyarilari"], ["Kararda genel kredi sözleşmesi tutarı yok; kullandırım kontrolünde 'Kontrol' bulgusu çıkar."])
        c = agent.tamamla(copy.deepcopy(CEVAP["ON-6101"]) | {"sozlesme_tutari": "20.000.000", "gecerlilik_bitis": ""})
        self.assertEqual(len(c["kod_uyarilari"]), 2)
        self.assertIn("limitin (25.000.000,00) altında", c["kod_uyarilari"][0])


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ornek_ve_duzeltme(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            self.assertEqual(agent.main(["--evet", "--cikti", str(cikti)]), 0)
            self.assertEqual(len(MESAJLAR), 3)                          # 2 karar + ON-6101 düzeltmesi
            for m in MESAJLAR:
                self.assertNotIn("Demirtaş", m)                            # kefil adları gönderilmez
            o = load_workbook(Path(tmp) / "t_onaylar.xlsx").active
            self.assertEqual([c.value for c in o[1]], agent.ONAY_BAS)
            self.assertEqual([c.value for c in o[2]][:8], ["ON-6101", "M60101", 25000000, "TL", "Rotatif, Ticari Taksitli", 36, 41, datetime(2027, 9, 30)])
            s = [[c.value for c in r] for r in load_workbook(Path(tmp) / "t_sartlar.xlsx").active.iter_rows(min_row=2)]
            self.assertEqual(len(s), 8)
            self.assertEqual(s[3][1:5], ["Müşteri çeki temliki (rotatif kullandırımlarda, limitin %20'si)", "Teminat", "Kullandırım öncesi", 5000000])
            self.assertIn("Hasan Demirtaş", s[1][1])                         # ad geri açılır
            self.assertEqual(s[4][8], datetime(2027, 3, 31))
            self.assertEqual(load_workbook(cikti)["Kontrol"]["A2"].value, "Son kontrol: Kod denetiminden sorunsuz geçti")

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(agent.main(["--kararlar", str(Path(tmp) / "yok.txt"), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
