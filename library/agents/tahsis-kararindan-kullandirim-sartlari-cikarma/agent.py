"""
Tahsis Kararından Kullandırım Şartları Çıkarma — Gerisi Bende AI Agent
Bankacılık Operasyonları › Kredi Operasyon Uzmanı

Kredi komitesi / tahsis kararlarının serbest metninden kullandırım kontrolüne girecek onay bilgilerini ve şartları çıkarır.
1. Kod kararları okur (.txt/.docx/.pdf ya da klasör; "Karar No:" satırı yeni kararı başlatır); telefon, e-posta, kimlik
   numarası ve --gizle adları maskelenir.
2. Model limit, ürünler, azami vade, asgari faiz, geçerlilik, sözleşme tutarını ve her teminat / ön koşul / belge /
   taahhüt şartını (zamanlama, tutar, termin, kararda geçen cümlesi) çıkarır.
3. Kod denetler: tutar, oran ve tarihler kararda geçiyor mu; her şartın kaynak cümlesi kararda aynen var mı; teminat ya
   da şart içeren her cümle bir şarta alınmış mı; zamanlama cümleyle uyumlu mu. Sorun varsa yalnız o karar modele geri
   gönderilir.
4. Kod yüzdeyle verilen teminat tutarlarını hesaplar ve "Kredi Kullandırım Ön Kontrolü" kod bloğunun onay ve şart
   dosyalarını yazar.

Kullanım:
    python agent.py                                                       # örnek 2 karar
    python agent.py --kararlar kararlar/ --gizle "Ad Soyad"
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KOD_DOLGU = PatternFill("solid", fgColor="E8F0FE")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}
UZANTILAR = {".txt", ".md", ".docx", ".pdf"}
BASLIK = re.compile(r"^\s*(?:Karar|Onay|Tahsis)\s+No\s*:\s*(\S+)", re.I | re.M)
SART_IZI = re.compile(r"\b(ipotek|rehin|rehni|kefalet|temlik|teminat mektub|blokaj|sigorta|ibraz|taahhut|alinacak|saglanacak|tesis edil|"
                      r"kgf|garanti fonu|cek|senet)\w*")
SONRA = re.compile(r"\b(takiben|kullandirimdan sonra|kullandirim sonrasi|gun icinde|ay icinde|tarihine kadar)\b")
ONCE = re.compile(r"\b(kullandirim oncesi\w*|kullandirimdan once|edilmeden kullandirim yapilma\w*)\b")
YUZDE = re.compile(r"%\s*(\d+(?:[.,]\d+)?)|(\d+(?:[.,]\d+)?)\s*%")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def para(x) -> Decimal | None:
    s = re.sub(r"[^\d,.\-]", "", str(x or ""))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1 or re.fullmatch(r"\d{1,3}\.\d{3}", s):
        s = s.replace(".", "")
    try:
        return Decimal(s) if s not in ("", "-", ".") else None
    except InvalidOperation:
        return None


def sayilar(metin: str) -> set[Decimal]:
    temiz = re.sub(r"\b\d{1,2}\.\d{1,2}\.\d{4}\b|\b[A-Z]{1,4}-[\w-]+\b", " ", str(metin))
    return {v for x in re.findall(r"\d+(?:[.,]\d+)*", temiz) if (v := para(x)) is not None}


def tarihler(metin: str) -> set[str]:
    return set(re.findall(r"\b\d{1,2}\.\d{1,2}\.\d{4}\b", str(metin)))


def tarih_coz(x) -> date | None:
    m = re.search(r"\b(\d{1,2})\.(\d{1,2})\.(\d{4})\b", str(x or ""))
    try:
        return date(int(m[3]), int(m[2]), int(m[1])) if m else None
    except ValueError:
        return None


def tl(x: Decimal | None) -> str:
    return "" if x is None else f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


# ----------------------------------------------------------------------------
# Girdi
# ----------------------------------------------------------------------------

def takma_adlar(adlar: list[str]) -> dict[str, str]:
    return {a: f"[KİŞİ-{i}]" for i, a in enumerate(dict.fromkeys(x.strip() for x in adlar if x.strip()), 1)}


def maskele(metin: str, kisiler: dict[str, str]) -> str:
    for g, t in sorted(kisiler.items(), key=lambda x: -len(x[0])):
        metin = re.sub(re.escape(g), t, metin, flags=re.I)
    return llm.maskele(metin)


def kararlari_oku(yol: Path, kisiler: dict[str, str]) -> list[dict]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: karar dosyası / klasörü bulunamadı.")
    dosyalar = sorted(p for p in yol.iterdir() if p.suffix.lower() in UZANTILAR) if yol.is_dir() else [yol]
    out = []
    for d in dosyalar:
        metin = belge.metin_oku(d).replace("\r", "")
        basl = list(BASLIK.finditer(metin))
        parcalar = [(m[1], metin[m.start():basl[i + 1].start() if i + 1 < len(basl) else len(metin)]) for i, m in enumerate(basl)] or [(d.stem, metin)]
        for no, p in parcalar:
            if p.strip():
                out.append({"no": no, "metin": maskele(p.strip(), kisiler), "dosya": d.name})
    if not out:
        raise llm.LLMHatasi(f"{yol}: okunabilir karar metni yok.")
    tekrar = sorted({k["no"] for k in out if [x["no"] for x in out].count(k["no"]) > 1})
    if tekrar:
        raise llm.LLMHatasi(f"Aynı karar numarası birden çok kez geçiyor: {', '.join(tekrar)}")
    return out


def cumleler(metin: str) -> list[str]:
    out = []
    for s in metin.splitlines():
        out += [c.strip() for c in re.split(r"(?<=[.!?])\s+(?=[A-ZÇĞİÖŞÜ])", s) if c.strip()]
    return out


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

SART = {"type": "object", "properties": {
    "sart": S, "tur": {"type": "string", "enum": ["Teminat", "Ön koşul", "Belge", "Taahhüt"]},
    "zamanlama": {"type": "string", "enum": ["Kullandırım öncesi", "Kullandırım sonrası"]}, "gereken_tutar": S, "termin": S, "kaynak": S},
    "required": ["sart", "tur", "zamanlama", "gereken_tutar", "termin", "kaynak"], "additionalProperties": False}
SEMA = {"type": "object", "properties": {
    "musteri_no": S, "musteri": S, "limit": S, "para_birimi": S, "urunler": {"type": "array", "items": S}, "azami_vade_ay": S, "asgari_faiz": S,
    "gecerlilik_bitis": S, "sozlesme_tutari": S, "sartlar": {"type": "array", "items": SART}, "belirsizlikler": {"type": "array", "items": S}},
    "required": ["musteri_no", "musteri", "limit", "para_birimi", "urunler", "azami_vade_ay", "asgari_faiz", "gecerlilik_bitis", "sozlesme_tutari",
                 "sartlar", "belirsizlikler"],
    "additionalProperties": False}


def yuzde_tutarlari(kaynak: str, limit: Decimal | None) -> set[Decimal]:
    """'limitin %20'si', 'kredi tutarının %80'i' gibi ifadelerden kodun hesapladığı tutarlar."""
    if not limit or not re.search(r"\b(limit|kredi tutar)", katla(kaynak)):
        return set()
    return {(limit * para(a or b) / 100).quantize(Decimal("0.01"), ROUND_HALF_UP) for a, b in YUZDE.findall(kaynak)}


def denetle(k: dict, c: dict) -> list[str]:
    no, metin, sorun = k["no"], k["metin"], []
    km = katla(metin)
    limit = para(c["limit"])
    if not limit or limit <= 0:
        sorun.append(f"{no}: limit tutarı yok ya da okunamadı ('{c['limit']}')")
    izinli = sayilar(metin)
    alanlar = [c["limit"], c["azami_vade_ay"], c["asgari_faiz"], c["sozlesme_tutari"]] + [s["sart"] for s in c["sartlar"]]
    for s in c["sartlar"]:
        alanlar.append(s["gereken_tutar"])
        izinli |= yuzde_tutarlari(s["kaynak"], limit)
    fazla = sorted(set().union(*(sayilar(x) for x in alanlar)) - izinli, key=float)
    if fazla:
        sorun.append(f"{no}: kararda olmayan sayı: {', '.join(f'{x:g}' for x in fazla)}")
    fazla_t = sorted(tarihler(" ".join([c["gecerlilik_bitis"]] + [s["termin"] for s in c["sartlar"]])) - tarihler(metin))
    if fazla_t:
        sorun.append(f"{no}: kararda olmayan tarih: {', '.join(fazla_t)}")
    if c["musteri_no"] and katla(c["musteri_no"]) not in km:
        sorun.append(f"{no}: müşteri no '{c['musteri_no']}' kararda yok")
    for u in c["urunler"]:
        if katla(u) not in km:
            sorun.append(f"{no}: ürün '{u}' kararda bu adla geçmiyor")
    kaynaklar = []
    for i, s in enumerate(c["sartlar"], 1):
        ks = katla(s["kaynak"])
        ad = f"şart {i} ('{s['sart']}')"
        if len(ks.split()) < 4 or ks not in km:
            sorun.append(f"{no}: {ad}: kaynak cümle kararda aynen geçmiyor; kararın ilgili cümlesini kelimesi kelimesine yazın")
            continue
        kaynaklar.append(ks)
        if SONRA.search(ks) and not ONCE.search(ks) and s["zamanlama"] != "Kullandırım sonrası":
            sorun.append(f"{no}: {ad}: cümlede kullandırım sonrası süre var ('{SONRA.search(ks)[0]}'); zamanlama 'Kullandırım sonrası' olmalı")
        if s["zamanlama"] == "Kullandırım sonrası" and not SONRA.search(ks):
            sorun.append(f"{no}: {ad}: cümlede kullandırım sonrasına işaret eden ifade yok; 'Kullandırım öncesi' seçin ve belirsizliğe yazın")
        tutarlar = {x for x in sayilar(s["kaynak"]) if x >= 1000}
        if s["gereken_tutar"] and para(s["gereken_tutar"]) not in tutarlar:
            sorun.append(f"{no}: {ad}: gereken tutar '{s['gereken_tutar']}' kaynak cümlede yok" + ("; yüzdeli tutarı boş bırakın, kod hesaplar" if YUZDE.search(s["kaynak"]) else ""))
        if tutarlar and not s["gereken_tutar"] and s["tur"] == "Teminat":
            sorun.append(f"{no}: {ad}: kaynak cümlede tutar var ({', '.join(tl(x) for x in sorted(tutarlar))}) ama gereken tutar boş")
        if tarihler(s["kaynak"]) and not s["termin"]:
            sorun.append(f"{no}: {ad}: kaynak cümlede tarih var ({', '.join(sorted(tarihler(s['kaynak'])))}) ama termin boş")
    for cu in cumleler(metin):
        kc = katla(cu)
        if len(kc.split()) >= 5 and not cu.rstrip().endswith(":") and SART_IZI.search(kc) and not any(x in kc or kc in x for x in kaynaklar):
            sorun.append(f"{no}: şart içeren cümle şartlara alınmamış: \"{cu[:120]}\"")
    return sorun


def uret(kararlar: list[dict], duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    sonuc = dict(onceki or {})
    hedef = kararlar if duzeltme is None else [k for k in kararlar if any(s.startswith(k["no"] + ":") for s in duzeltme)]
    for k in hedef:
        mesaj = f"<karar no=\"{k['no']}\">\n{k['metin']}\n</karar>"
        d = [s.split(": ", 1)[1] for s in (duzeltme or []) if s.startswith(k["no"] + ":")]
        if d:
            mesaj += ("\n<duzeltme>\nÖnceki cevap: " + json.dumps(sonuc.get(k["no"]), ensure_ascii=False) + "\nKodun bulduğu sorunlar:\n"
                      + "\n".join(f"- {s}" for s in d) + "\n</duzeltme>")
        sonuc[k["no"]] = llm.json_iste(sistem, mesaj, SEMA)
    return sonuc


def duzeltme_dongusu(uret_f, denetle_f, tur: int) -> tuple[dict, list[str], int]:
    taslak = uret_f(None, None)
    kalan = denetle_f(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret_f(kalan, taslak)
        kalan = denetle_f(taslak)
    return taslak, kalan, yapilan


def geri_ac(s, kisiler: dict[str, str]):
    if isinstance(s, str):
        for g, t in kisiler.items():
            s = s.replace(t, g)
        return s
    if isinstance(s, list):
        return [geri_ac(x, kisiler) for x in s]
    if isinstance(s, dict):
        return {k: geri_ac(v, kisiler) for k, v in s.items()}
    return s


# ----------------------------------------------------------------------------
# Hesap ve çıktı
# ----------------------------------------------------------------------------

def tamamla(c: dict) -> dict:
    """Yüzdeyle verilen teminat tutarını hesaplar; kod uyarılarını ekler."""
    limit = para(c["limit"])
    uyarilar = []
    for s in c["sartlar"]:
        s["hesap"] = ""
        t = yuzde_tutarlari(s["kaynak"], limit)
        if not s["gereken_tutar"] and len(t) == 1:
            s["gereken_tutar"] = tl(next(iter(t)))
            s["hesap"] = f"Kod hesapladı: limit {tl(limit)} × %{next(a or b for a, b in YUZDE.findall(s['kaynak']))}"
    soz = para(c["sozlesme_tutari"])
    if not c["sozlesme_tutari"]:
        uyarilar.append("Kararda genel kredi sözleşmesi tutarı yok; kullandırım kontrolünde 'Kontrol' bulgusu çıkar.")
    elif soz is not None and limit and soz < limit:
        uyarilar.append(f"Sözleşme tutarı ({tl(soz)}) limitin ({tl(limit)}) altında; limitin tamamı kullandırılamaz.")
    b = tarih_coz(c["gecerlilik_bitis"])
    if not b:
        uyarilar.append("Limit geçerlilik tarihi kararda bulunamadı.")
    c["kod_uyarilari"] = uyarilar
    return c


def tablo(ws, bas: list[str], satirlar: list[list], ai: set[int] = frozenset(), kod: set[int] = frozenset(), genis: dict | None = None) -> None:
    ws.append(bas)
    for h in ws[1]:
        h.fill, h.font, h.alignment = BASLIK_DOLGU, BASLIK_YAZI, UST
    for r in satirlar:
        ws.append(r)
        n = ws.max_row
        for j, v in enumerate(r, 1):
            ws.cell(n, j).alignment = UST
            if isinstance(v, date):
                ws.cell(n, j).number_format = "DD.MM.YYYY"
            elif isinstance(v, float):
                ws.cell(n, j).number_format = "#,##0.00"
            if j in ai and v not in (None, ""):
                ws.cell(n, j).fill = AI_DOLGU
            if j in kod and v not in (None, ""):
                ws.cell(n, j).fill = KOD_DOLGU
    for j in range(1, len(bas) + 1):
        ws.column_dimensions[get_column_letter(j)].width = (genis or {}).get(j, 16)
    ws.freeze_panes = "A2"


def sayi(x) -> float | None:
    v = para(x)
    return float(v) if v is not None else None


ONAY_BAS = ["Onay No", "Müşteri No", "Onaylı Limit", "Para Birimi", "Ürünler", "Azami Vade (ay)", "Asgari Faiz (%)", "Geçerlilik Bitiş", "Mevcut Risk",
            "Sözleşme Tutarı", "Sözleşme Durumu"]
SART_BAS = ["Onay No", "Şart", "Tür", "Zamanlama", "Gereken Tutar", "Tesis Edilen Tutar", "Durum", "Tamamlanma Tarihi", "Termin"]


def satirlar(kararlar: list[dict], sonuc: dict) -> tuple[list[list], list[list]]:
    o, s = [], []
    for k in kararlar:
        c = sonuc[k["no"]]
        o.append([k["no"], c["musteri_no"], sayi(c["limit"]), (c["para_birimi"] or "TL").upper().replace("TRY", "TL"), ", ".join(c["urunler"]) or "*",
                  sayi(c["azami_vade_ay"]), sayi(c["asgari_faiz"]), tarih_coz(c["gecerlilik_bitis"]), None, sayi(c["sozlesme_tutari"]), ""])
        for x in c["sartlar"]:
            s.append([k["no"], x["sart"], x["tur"], x["zamanlama"], sayi(x["gereken_tutar"]), None, "Bekliyor", None, tarih_coz(x["termin"])])
    return o, s


def rapor_yaz(cikti: Path, kararlar: list[dict], sonuc: dict, kalan: list[str], tur: int) -> tuple[Path, Path]:
    o, s = satirlar(kararlar, sonuc)
    wb = Workbook()
    ws = wb.active
    ws.title = "Onaylar"
    tablo(ws, ONAY_BAS + ["Müşteri", "Kaynak Dosya", "Son Kontrol"],
          [r + [sonuc[k["no"]]["musteri"], k["dosya"], "; ".join(x.split(": ", 1)[1] for x in kalan if x.startswith(k["no"] + ":")) or "Geçti"]
           for r, k in zip(o, kararlar)], ai=set(range(2, 9)) | {10, 12}, genis={5: 26, 12: 30, 14: 50})
    for i in range(2, ws.max_row + 1):
        ws.cell(i, 14).fill = YESIL if ws.cell(i, 14).value == "Geçti" else KIRMIZI
    ws = wb.create_sheet("Şartlar")
    ek = [[x["kaynak"], x["hesap"]] for k in kararlar for x in sonuc[k["no"]]["sartlar"]]
    tablo(ws, SART_BAS + ["Kararda Geçen Cümle", "Kod Hesabı"], [r + e for r, e in zip(s, ek)], ai={2, 3, 4, 5, 9}, kod={11}, genis={2: 34, 10: 60, 11: 34})
    for i, (_, e) in enumerate(zip(s, ek), 2):
        if e[1]:
            ws.cell(i, 5).fill = KOD_DOLGU
    ws = wb.create_sheet("Açık Noktalar")
    tablo(ws, ["Onay No", "Kaynak", "Açıklama"],
          [[k["no"], "Model", b] for k in kararlar for b in sonuc[k["no"]]["belirsizlikler"]] + [[k["no"], "Kod", u] for k in kararlar for u in sonuc[k["no"]]["kod_uyarilari"]],
          genis={3: 90})
    k = wb.create_sheet("Kontrol")
    k.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    k.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    k.cell(2, 1).fill = KIRMIZI if kalan else YESIL
    k.append(["Mor hücreler modelin çıkardığı, mavi hücreler kodun hesapladığı değerlerdir. Her satırı karar aslıyla karşılaştırın. Mevcut Risk, "
              "Sözleşme Durumu, Tesis Edilen Tutar, Durum ve Tamamlanma Tarihi sütunlarını kendi kayıtlarınızdan doldurup onay ve şart dosyalarını "
              "'Kredi Kullandırım Ön Kontrolü' kod bloğuna verin."])
    k.column_dimensions["A"].width = 140
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    yollar = []
    for ad, bas, r in (("onaylar", ONAY_BAS, o), ("sartlar", SART_BAS, s)):
        w = Workbook()
        tablo(w.active, bas, r)
        y = cikti.with_name(f"{cikti.stem}_{ad}.xlsx")
        w.save(y)
        yollar.append(y)
    return yollar[0], yollar[1]


def calistir(yol: Path, cikti: Path, gizle: list[str] | None = None, tur: int = 2, evet: bool = False) -> dict:
    kisiler = takma_adlar(gizle or [])
    kararlar = kararlari_oku(yol, kisiler)
    print(f"[OK] {len(kararlar)} tahsis kararı okundu")
    llm.onay_al(f"{len(kararlar)} tahsis kararının metni gönderilecek (telefon, e-posta, TCKN, IBAN maskeli"
                f"{f'; {len(kisiler)} ad takma adla' if kisiler else ''}).", evet)
    sonuc, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(kararlar, d, o), lambda x: [s for k in kararlar for s in denetle(k, x[k["no"]])], tur)
    sonuc = {n: tamamla(c) for n, c in geri_ac(sonuc, kisiler).items()}
    onay, sart = rapor_yaz(cikti, kararlar, sonuc, kalan, yapilan)
    return {"kararlar": kararlar, "sonuc": sonuc, "kalan": kalan, "tur": yapilan, "onaylar": onay, "sartlar": sart}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Tahsis kararlarından kullandırım kontrolüne girecek onay bilgilerini ve şartları çıkarır.")
    p.add_argument("--kararlar", type=Path, default=ORNEK / "tahsis_kararlari.txt", help="Karar dosyası (.txt/.docx/.pdf) ya da klasör; 'Karar No:' satırı yeni kararı başlatır")
    p.add_argument("--gizle", nargs="*", default=[], metavar="AD", help="Kararda geçen ve gönderilmeyecek kişi adları (kefiller, ortaklar)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "tahsis_sartlari.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    gizle = a.gizle or (["Hasan Demirtaş", "Selin Demirtaş"] if a.kararlar == ORNEK / "tahsis_kararlari.txt" else [])
    try:
        s = calistir(a.kararlar, a.cikti, gizle, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    for k in s["kararlar"]:
        c = s["sonuc"][k["no"]]
        on = sum(x["zamanlama"] == "Kullandırım öncesi" for x in c["sartlar"])
        print(f"[OK] {k['no']} · limit {c['limit']} {c['para_birimi']} · {len(c['sartlar'])} şart ({on} kullandırım öncesi) · "
              f"{len(c['belirsizlikler']) + len(c['kod_uyarilari'])} açık nokta")
    for x in s["kalan"]:
        print(f"[!] Kalan: {x}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[OK] Kod bloğu girdileri: {s['onaylar'].name}, {s['sartlar'].name}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
