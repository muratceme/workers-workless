"""
Acil Durum Eylem Planı Taslağı — Gerisi Bende AI Agent
İş Sağlığı ve Güvenliği › İş Güvenliği Uzmanı

1. Kod işyeri bilgi dosyasını okur: unvan, adres, tehlike sınıfı, vardiyalar, toplanma yeri, özel durumlar; bölümler
   (B1, B2 …), risk değerlendirmesi sonuçları (R1, R2 …) ve acil durum ekipmanları (E1, E2 …).
2. Kod İşyerlerinde Acil Durumlar Hakkında Yönetmelik'e göre hesaplar:
   - Söndürme, kurtarma ve koruma ekiplerinin her biri için destek elemanı sayısı (çok tehlikeli / tehlikeli / az tehlikeli:
     her 30 / 40 / 50 çalışana kadar en az 1; 10'dan az çalışanda üç ekip için tek kişi yeterli), vardiya bazında
   - İlk yardım ekibi (İlkyardım Yönetmeliği: her 10 / 15 / 20 çalışana kadar 1 ilk yardımcı)
   - Planın geçerlilik tarihi (en geç 2 / 4 / 6 yılda bir yenilenir), tatbikat son tarihi (en geç yılda bir)
   - Planda bulunması gereken bilgilerden işyeri dosyasında eksik olanlar
3. Model acil durum senaryolarını, önleyici tedbirleri, ekiplere göre müdahale adımlarını ve tahliye adımlarını yazar.
4. Kod denetler:
   - Her risk (R) en az bir senaryoda ele alınmış mı; yangın ve doğal afet senaryosu var mı
   - Yangın senaryosunda söndürme ekibi, doğal afet senaryosunda kurtarma ekibi görevli mi
   - Tahliyede toplanma yeri ve sayım var mı; özel durumlarda (gebe, engelli, kreş) refakat, ziyaretçi varsa yönlendirme var mı
   - İşyeri bilgisinde olmayan sayı (ekip sayısı, kapasite, mesafe) ya da eski acil numara (110, 155, 156) yazılmış mı
   Sorun varsa modele geri gönderilir (en fazla --tur kez).
5. Kod plan metnini, ekip tablosunu ve doküman kontrol listesini oluşturur.

Kullanım:
    python agent.py                                    # örnek (kurgusal) plastik ambalaj fabrikası
    python agent.py --girdi isyeri.docx --tur 3
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
ALANLAR = {"unvan": "Unvan", "adres": "Adres", "isveren": "İşveren", "hazirlayanlar": "Hazırlayanlar", "tarih": "Hazırlama Tarihi",
           "sinif": "Tehlike Sınıfı", "vardiyalar": "Vardiyalar", "calisan": "Çalışan Sayısı", "toplanma": "Toplanma Yeri",
           "ozel": "Özel Durumlar"}
BOLUMLER = {"bolumler": "B", "riskler": "R", "ekipman": "E"}
# Acil Durumlar Yönetmeliği md. 11/3 (2021 değişikliği), İlkyardım Yönetmeliği, md. 14/2 yenileme süresi
SINIF = {"cok tehlikeli": ("Çok tehlikeli", 30, 10, 2), "tehlikeli": ("Tehlikeli", 40, 15, 4), "az tehlikeli": ("Az tehlikeli", 50, 20, 6)}
EKIPLER = ["Söndürme", "Kurtarma", "Koruma", "İlk yardım"]
TURLER = ["Yangın", "Patlama", "Kimyasal yayılım", "Doğal afet", "Sabotaj", "Salgın hastalık", "Diğer"]
ESKI_NUMARA = {"110", "155", "156", "177"}
REFAKAT = re.compile(r"gebe|hamile|engelli|tekerlekli|yasli|kres|cocuk")
ZIYARETCI = re.compile(r"ziyaretci|musteri|tedarikci|sofor|misafir")
# Md. 12: planda (tahliye planı / kroki) bulunması gereken ve işyeri dosyasından aranan bilgiler
DOKUMAN = [("Yangın söndürme ekipmanlarının yerleri", r"sondurucu|yangin dolabi|sprinkler"), ("İlk yardım malzemelerinin yerleri", r"ilk ?yardim"),
           ("Elektrik kesim noktası", r"elektrik|salter"), ("Gaz akışının kesim noktası / vanası", r"vana|dogalgaz|lpg|gaz"),
           ("Uyarı (alarm) sistemi", r"alarm|anons|siren")]

SEMA = {
    "type": "object",
    "properties": {
        "senaryolar": {"type": "array", "items": {"type": "object", "properties": {
            "ad": S, "tur": {"type": "string", "enum": TURLER}, "dayanak": {"type": "array", "items": S},
            "onleyici_tedbirler": {"type": "array", "items": S},
            "mudahale": {"type": "array", "items": {"type": "object", "properties": {
                "ekip": {"type": "string", "enum": EKIPLER + ["Koordinasyon sorumlusu", "Tüm çalışanlar"]}, "adim": S},
                "required": ["ekip", "adim"], "additionalProperties": False}}},
            "required": ["ad", "tur", "dayanak", "onleyici_tedbirler", "mudahale"], "additionalProperties": False}},
        "tahliye": {"type": "array", "items": S},
    },
    "required": ["senaryolar", "tahliye"],
    "additionalProperties": False,
}


def isyeri_oku(yol: Path) -> dict:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: işyeri bilgi dosyası bulunamadı.")
    n = {a: "" for a in ALANLAR} | {b: [] for b in BOLUMLER}
    bolum = None
    for s in belge.metin_oku(yol).splitlines():
        if not s.strip():
            continue
        a, _, d = s.partition(":")
        anahtar = next((k for k, v in ALANLAR.items() if katla(v) == katla(a)), None)
        if anahtar and d.strip() and not s.lstrip().startswith(("-", "•", "*")):
            n[anahtar], bolum = d.strip(), None
        elif katla(s) in {katla(x) for x in ("Bölümler", "Riskler", "Ekipman")}:
            bolum = next(b for b in BOLUMLER if katla(b)[:5] == katla(s)[:5])
        elif bolum:
            n[bolum].append(s.strip().lstrip("-•* ").strip())
    if katla(n["sinif"]) not in SINIF:
        raise llm.LLMHatasi(f"{yol.name}: 'Tehlike Sınıfı:' satırı 'Az tehlikeli', 'Tehlikeli' ya da 'Çok tehlikeli' olmalı.")
    if not n["riskler"]:
        raise llm.LLMHatasi(f"{yol.name}: 'Riskler' bölümü gerekli (risk değerlendirmesi sonuçları).")
    n["vardiya"] = [(m.group(1).strip(), int(m.group(2))) for m in re.finditer(r"([^;,\d]+?)\s*(\d+)", n["vardiyalar"])]
    if not n["vardiya"]:
        if not str(n["calisan"]).strip().isdigit():
            raise llm.LLMHatasi(f"{yol.name}: 'Vardiyalar: Gündüz 85; Gece 32' ya da 'Çalışan Sayısı: 40' satırı gerekli.")
        n["vardiya"] = [("Tek vardiya", int(n["calisan"]))]
    n["toplam"] = sum(x for _, x in n["vardiya"])
    n["tarih_v"] = tarih_coz(n["tarih"]) or date.today()
    return n


def kimlikler(n: dict) -> dict:
    return {f"{h}{i}": s for b, h in BOLUMLER.items() for i, s in enumerate(n[b], 1)}


def ekip_hesapla(n: dict) -> dict:
    ad, kisi, ilk_yardim, yil = SINIF[katla(n["sinif"])]
    satirlar = []
    for vardiya, sayi in n["vardiya"]:
        if n["toplam"] < 10:
            s = {"Söndürme + Kurtarma + Koruma (tek kişi)": 1}
        else:
            s = {e: math.ceil(sayi / kisi) for e in EKIPLER[:3]}
        s["İlk yardım"] = math.ceil(sayi / ilk_yardim)
        satirlar.append({"vardiya": vardiya, "calisan": sayi, "ekipler": s})
    gecerlilik = n["tarih_v"].replace(year=n["tarih_v"].year + yil) if not (n["tarih_v"].month == 2 and n["tarih_v"].day == 29) else date(n["tarih_v"].year + yil, 2, 28)
    tatbikat = n["tarih_v"].replace(year=n["tarih_v"].year + 1) if not (n["tarih_v"].month == 2 and n["tarih_v"].day == 29) else date(n["tarih_v"].year + 1, 2, 28)
    return {"sinif": ad, "kisi": kisi, "ilk_yardim": ilk_yardim, "yil": yil, "vardiyalar": satirlar, "gecerlilik": gecerlilik, "tatbikat": tatbikat}


def dokuman_eksikleri(n: dict) -> list[str]:
    ekip = katla(" ".join(n["ekipman"]))
    eksik = [ad for ad, desen in DOKUMAN if not re.search(desen, ekip)]
    for alan, ad in (("toplanma", "Toplanma yeri"), ("adres", "İşyeri adresi"), ("hazirlayanlar", "Hazırlayanlar")):
        if not n[alan]:
            eksik.append(ad)
    return eksik


def denetle(t: dict, n: dict) -> list[str]:
    sorun = []
    k = kimlikler(n)
    izinli = sayilar(" ".join(k.values()) + " " + n["toplanma"] + " " + n["ozel"]) | {"112"}
    kullanilan = set()
    for i, s in enumerate(t["senaryolar"], 1):
        ad = f"S{i} ({s['ad'][:40]})"
        gec = [d for d in s["dayanak"] if d not in k]
        if not s["dayanak"] or gec:
            sorun.append(f"{ad}: dayanak yok ya da geçersiz {gec} (R, B, E numaraları kullanın)")
        kullanilan |= set(s["dayanak"])
        if not s["onleyici_tedbirler"]:
            sorun.append(f"{ad}: önleyici / sınırlandırıcı tedbir yazılmamış")
        ekipler = {m["ekip"] for m in s["mudahale"]}
        if not s["mudahale"]:
            sorun.append(f"{ad}: müdahale adımı yok")
        elif s["tur"] == "Yangın" and "Söndürme" not in ekipler:
            sorun.append(f"{ad}: yangın senaryosunda söndürme ekibine görev verilmemiş")
        elif s["tur"] == "Doğal afet" and "Kurtarma" not in ekipler:
            sorun.append(f"{ad}: doğal afet senaryosunda kurtarma ekibine görev verilmemiş")
        metin = " ".join([s["ad"]] + s["onleyici_tedbirler"] + [m["adim"] for m in s["mudahale"]])
        sorun += _sayi_sorunlari(ad, metin, izinli)
    for r in (x for x in k if x.startswith("R")):
        if r not in kullanilan:
            sorun.append(f"{r}: risk hiçbir senaryoda ele alınmamış ('{k[r][:60]}')")
    turler = {s["tur"] for s in t["senaryolar"]}
    for tur in ("Yangın", "Doğal afet"):
        if tur not in turler:
            sorun.append(f"Senaryolar: '{tur}' senaryosu yok (yönetmelik md. 8 bu ihtimalin değerlendirilmesini ister)")
    th = katla(" ".join(t["tahliye"]))
    if not t["tahliye"]:
        sorun.append("Tahliye: adım yok")
    else:
        if "toplanma" not in th:
            sorun.append("Tahliye: toplanma yerine gidiş adımı yok")
        if "sayim" not in th and "sayil" not in th:
            sorun.append("Tahliye: toplanma yerinde sayım adımı yok (yönetmelik md. 10/2)")
        if REFAKAT.search(katla(n["ozel"])) and "refakat" not in th:
            sorun.append("Tahliye: gebe / engelli çalışanlar için refakat adımı yok (yönetmelik md. 10/4)")
        if ZIYARETCI.search(katla(n["ozel"])) and not ZIYARETCI.search(th):
            sorun.append("Tahliye: ziyaretçi ve tedarikçilerin yönlendirilmesi yazılmamış (yönetmelik md. 10/6)")
        sorun += _sayi_sorunlari("Tahliye", " ".join(t["tahliye"]), izinli)
    return sorun


def _sayi_sorunlari(ad: str, metin: str, izinli: set) -> list[str]:
    sorun = []
    metin = re.sub(r"\b[BRE]\d+\b", " ", metin)  # "(E5)", "R1" gibi kimlikler sayı sayılmaz

    eski = sorted(ESKI_NUMARA & sayilar(metin))
    if eski:
        sorun.append(f"{ad}: eski acil numarası ({', '.join(eski)}); tüm acil çağrılar için 112 yazın")
    fazla = sorted(sayilar(metin) - izinli - ESKI_NUMARA)
    if fazla:
        sorun.append(f"{ad}: işyeri bilgisinde olmayan sayı: {', '.join(fazla)} (ekip sayılarını kod hesaplar; sayı uydurmayın)")
    return sorun


def uret(n: dict, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    satirlar = "\n".join(f"{i} | {s}" for i, s in kimlikler(n).items())
    mesaj = (f"<tehlike_sinifi>{n['sinif']}</tehlike_sinifi>\n<toplanma_yeri>{n['toplanma'] or '(belirtilmemiş)'}</toplanma_yeri>\n"
             f"<ozel_durumlar>{n['ozel'] or '(yok)'}</ozel_durumlar>\n<isyeri>\n{satirlar}\n</isyeri>" + duzeltme_blogu(duzeltme, onceki))
    return llm.json_iste(sistem, mesaj, SEMA)


def plan_metni(t: dict, n: dict, e: dict) -> str:
    s = ["ACİL DURUM PLANI (TASLAK)", "", f"İşyeri unvanı: {n['unvan']}", f"Adres: {n['adres']}", f"İşveren: {n['isveren']}",
         f"Hazırlayanlar: {n['hazirlayanlar']}", f"Hazırlandığı tarih: {n['tarih_v']:%d.%m.%Y}",
         f"Geçerlilik tarihi: {e['gecerlilik']:%d.%m.%Y} ({e['sinif'].lower()} sınıf: en geç {e['yil']} yılda bir yenilenir)",
         f"Tehlike sınıfı: {e['sinif']} · Çalışan: {n['toplam']}", "", "1. BELİRLENEN ACİL DURUMLAR VE MÜDAHALE"]
    for i, x in enumerate(t["senaryolar"], 1):
        s += ["", f"1.{i} {x['ad']} ({x['tur']})", "Önleyici ve sınırlandırıcı tedbirler:"] + [f"  - {y}" for y in x["onleyici_tedbirler"]]
        s += ["Müdahale:"] + [f"  {j}. [{m['ekip']}] {m['adim']}" for j, m in enumerate(x["mudahale"], 1)]
    s += ["", "2. TAHLİYE", f"Toplanma yeri: {n['toplanma']}"] + [f"  {j}. {y}" for j, y in enumerate(t["tahliye"], 1)]
    s += ["", "3. ACİL DURUM EKİPLERİ (destek elemanı sayısı en az)"]
    for v in e["vardiyalar"]:
        s.append(f"  {v['vardiya']} ({v['calisan']} çalışan): " + ", ".join(f"{k} {x}" for k, x in v["ekipler"].items()))
    s += ["  Her ekipte bir ekip başı bulunur; koordinasyon için koruma ekibinden sorumlu görevlendirilir. Ad, soyad, unvan, sorumluluk alanı ve",
          "  iletişim bilgilerini içeren ekip listesi görünür biçimde asılır.", "", "4. ACİL DURUM İRTİBAT",
          "  Acil çağrı (itfaiye, ambulans, polis, AFAD): 112", "", "5. TATBİKAT",
          f"  İlk tatbikat en geç {e['tatbikat']:%d.%m.%Y}; sonrasında en geç yılda bir. Tatbikat formu doldurulur, eksikler giderilir.",
          "", "Sayfalar numaralandırılıp hazırlayanlarca paraflanır, son sayfa imzalanır; tahliye planı (kroki) giriş-çıkışlarda ve katlarda asılır."]
    return "\n".join(s)


def rapor_yaz(cikti: Path, t, n, e, eksik, kalan, tur, metin) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Senaryolar"
    k = kimlikler(n)
    baslik(ws, ["No", "Senaryo", "Tür", "Dayanak", "Önleyici Tedbirler", "Müdahale Adımları"], (6, 30, 14, 40, 50, 70))
    for i, x in enumerate(t["senaryolar"], 1):
        ws.append([f"S{i}", x["ad"], x["tur"], "\n".join(f"{d}. {k.get(d, '?')}" for d in x["dayanak"]), "\n".join(x["onleyici_tedbirler"]),
                   "\n".join(f"[{m['ekip']}] {m['adim']}" for m in x["mudahale"])])
        boya(ws, (2, 5, 6), AI_DOLGU)
    ws.append([])
    ws.append(["", "Tahliye", "", "", "", "\n".join(f"{j}. {y}" for j, y in enumerate(t["tahliye"], 1))])
    boya(ws, (6,), AI_DOLGU)
    son_not(ws, "Önemli: Taslaktır. İş güvenliği uzmanı, işyeri hekimi ve işveren tarafından işyerinde doğrulanmalı; kroki ve ekip listesi eklenmelidir.", kalan, tur)
    hizala(ws)
    w = wb.create_sheet("Ekipler")
    baslik(w, ["Vardiya", "Çalışan", "Ekip", "En Az Destek Elemanı", "Ad Soyad / Unvan / İletişim (doldurulacak)"], (16, 10, 40, 20, 50))
    for v in e["vardiyalar"]:
        for ekip, sayi in v["ekipler"].items():
            w.append([v["vardiya"], v["calisan"], ekip, sayi, ""])
    w.append([])
    w.append([f"Dayanak: Acil Durumlar Yönetmeliği md. 11 ({e['sinif'].lower()}: her {e['kisi']} çalışana kadar 1); İlkyardım Yönetmeliği (her "
              f"{e['ilk_yardim']} çalışana kadar 1 ilk yardımcı). Vardiya düzeni dikkate alınmıştır; yedek görevlendirin."])
    hizala(w)
    d = wb.create_sheet("Doküman Kontrol")
    baslik(d, ["Plan Unsuru (md. 12)", "Durum"], (50, 40))
    for ad in ("Unvan, adres, işveren", "Hazırlayanlar", "Hazırlama ve geçerlilik tarihi", "Belirlenen acil durumlar", "Önleyici tedbirler",
               "Müdahale ve tahliye yöntemleri"):
        d.append([ad, "Taslakta var" if ad != "Hazırlayanlar" or n["hazirlayanlar"] else "Eksik"])
    for ad, _ in DOKUMAN:
        d.append([ad, "İşyeri bilgisinde eksik: krokiye ekleyin" if ad in eksik else "İşyeri bilgisinde var: krokide gösterin"])
        boya(d, (2,), SARI if ad in eksik else YESIL)
    for ad in ("Kaçış yolları ve toplanma yeri (kroki)", "Özel risk bölümleri (kimyasal, patlama) (kroki)", "Ekip listesi (ad, unvan, iletişim)"):
        d.append([ad, "Elle eklenecek"])
        boya(d, (2,), SARI)
    hizala(d)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    cikti.with_suffix(".txt").write_text(metin, encoding="utf-8")


def calistir(girdi: Path, cikti: Path, tur: int = 2, evet: bool = False) -> dict:
    n = isyeri_oku(girdi)
    e = ekip_hesapla(n)
    eksik = dokuman_eksikleri(n)
    print(f"[OK] {n['unvan']} · {e['sinif']} · {n['toplam']} çalışan · {len(n['riskler'])} risk")
    for x in eksik:
        print(f"[!] İşyeri bilgisinde eksik: {x}")
    llm.onay_al(f"Bölüm, risk ve ekipman bilgileri ({len(kimlikler(n))} satır) gönderilecek; unvan, adres ve hazırlayanlar gönderilmez.", evet)
    t, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(n, d, o), lambda x: denetle(x, n), tur)
    metin = plan_metni(t, n, e)
    rapor_yaz(cikti, t, n, e, eksik, kalan, yapilan, metin)
    return {"isyeri": n, "ekip": e, "eksik": eksik, "taslak": t, "kalan": kalan, "tur": yapilan, "metin": metin}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    p = argparse.ArgumentParser(description="İşyeri bilgilerinden acil durum senaryoları, ekipler ve tahliye adımlarını içeren plan taslağı hazırlar.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "isyeri.txt",
                   help="İşyeri bilgisi (.txt/.docx): Unvan, Adres, Tehlike Sınıfı, Vardiyalar, Toplanma Yeri satırları; Bölümler, Riskler, Ekipman bölümleri")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "acil_durum_plani.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.cikti, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(s['taslak']['senaryolar'])} senaryo · geçerlilik {s['ekip']['gecerlilik']:%d.%m.%Y}")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Plan: {a.cikti.with_suffix('.txt').resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
