"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import json
import os
import re
import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
YAS = ORNEK / "prim_alacak_yaslandirma.xlsx"
KAD = ORNEK / "kademeler.csv"
BUGUN = date(2026, 10, 1)
CEVAP = json.loads((Path(__file__).parent / "sahte_cevaplar.json").read_text(encoding="utf-8"))
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    """Sigortalıya son hatırlatmada ilk turda 'teminatınız durabilir' yazar; düzeltmede kaldırır."""
    MESAJLAR.append(kullanici)
    if "Alıcı türü: acente" in kullanici:
        ilk = re.search(r"Poliçe (P-\d+)", kullanici)[1]
        acente = {"P-1": "Örnek Acente A", "P-5": "Örnek Acente C", "P-3": "Örnek Acente B"}[ilk]
        return copy.deepcopy(CEVAP["acente"][acente])
    p = re.search(r"Poliçe (P-\d+)", kullanici)[1]
    return copy.deepcopy(CEVAP["sigortali"]["P-2_hatali" if p == "P-2" and "<duzeltme>" not in kullanici else p])


def mesajlar(alici="sigortali", odeme="", imza="Tahsilat Birimi", sinir=agent.SMS_SINIR):
    m = agent.gruplar(agent.taksitleri_oku(YAS, BUGUN), alici, agent.kademeleri_oku(KAD))
    for x in m:
        x["sms_sinir"] = sinir
        x["butce"] = sinir - len(agent.tam_metin(x, {"govde": ""}, "sms", odeme, imza))
    return m


class KodTesti(unittest.TestCase):
    def test_okuma_ve_gruplama(self):
        t = agent.taksitleri_oku(YAS, BUGUN)
        self.assertEqual(len(t), 10)
        k = agent.kademeleri_oku(KAD)
        self.assertEqual([(x["ad"], x["gun"], bool(x["zorunlu"])) for x in k], [("Nazik hatırlatma", 1, False), ("İkinci hatırlatma", 16, False), ("Son hatırlatma", 46, True)])
        m = mesajlar()
        self.assertEqual([(x["anahtar"], x["kademe"]["ad"], x["toplam"], x["en_uzun"]) for x in m],
                         [("P-2", "Son hatırlatma", D(90000), 108), ("P-5", "İkinci hatırlatma", D(30000), 31), ("P-1", "İkinci hatırlatma", D(30000), 30),
                          ("P-4", "İkinci hatırlatma", D(12000), 30), ("P-3", "Nazik hatırlatma", D(80000), 6)])
        a = mesajlar("acente")
        self.assertEqual([(x["anahtar"], x["toplam"], [y["police"] for y in x["taksitler"]]) for x in a],
                         [("Örnek Acente A", D(120000), ["P-1", "P-2", "P-2", "P-2"]), ("Örnek Acente C", D(30000), ["P-5"]),
                          ("Örnek Acente B", D(92000), ["P-3", "P-4"])])
        self.assertIn("sigortalı: [SİGORTALI-2]", agent.veri_metni(a[0], "acente"))
        self.assertNotIn("Kurgu Mobilya", agent.veri_metni(m[0], "sigortali") + agent.veri_metni(a[0], "acente"))
        with tempfile.TemporaryDirectory() as tmp:                      # gecikme sütunu yoksa vade ve tarihten hesaplanır
            y = Path(tmp) / "t.csv"
            y.write_text("Poliçe No;Vade Tarihi;Taksit Tutarı;Tahsil Edilen\nX-1;01.09.2026;1.000,00;400,00\nX-2;01.09.2026;500;500\n", encoding="utf-8")
            self.assertEqual([(x["police"], x["kalan"], x["gecikme"]) for x in agent.taksitleri_oku(y, BUGUN)], [("X-1", D("600.00"), 30)])

    def test_denetim(self):
        m = {x["anahtar"]: x for x in mesajlar()}
        self.assertEqual([agent.denetle(m[p], CEVAP["sigortali"][p], "eposta") for p in ("P-2", "P-5", "P-1", "P-4", "P-3")], [[]] * 5)
        self.assertIn("'teminatınız durabilir'", agent.denetle(m["P-2"], CEVAP["sigortali"]["P-2_hatali"], "eposta")[0])
        c = {"konu": "", "govde": "Sayın müşterimiz, P-2 poliçenizin borcu 95.000,00 TL; 15.06.2026 vadeli. Aksi hâlde icra takibi başlatılacaktır. [SİGORTALI-1]"}
        s = agent.denetle(m["P-2"], c, "eposta")
        for parca in ("toplam tutar '90.000,00'", "vade '15.07.2026'", "veride olmayan sayı / tarih: 95000", "'icra'", "hitapla", "yer tutucusu", "konusu boş"):
            self.assertTrue(any(parca in x for x in s), parca)
        a = {x["anahtar"]: x for x in mesajlar("acente")}
        self.assertEqual([agent.denetle(a[k], v, "sms") for k, v in CEVAP["acente"].items()], [[]] * 3)
        uzun = {"konu": "", "govde": CEVAP["acente"]["Örnek Acente A"]["govde"] + " " + "x" * 200}
        self.assertIn("en fazla 282 karakter", agent.denetle(a["Örnek Acente A"], uzun, "sms")[-1])

    def test_tam_metin(self):
        m = mesajlar()
        t = agent.tam_metin(m[0], CEVAP["sigortali"]["P-2"], "eposta", "IBAN TR..", "Tahsilat")
        self.assertTrue(t.startswith("Sayın Kurgu Mobilya Ltd. Şti.,\n\nP-2 numaralı"))
        self.assertIn("genel şartları ve ilgili mevzuat", t)                    # son hatırlatmanın zorunlu cümlesi
        self.assertTrue(t.endswith("Ödeme bilgisi: IBAN TR..\n\nSaygılarımızla,\nTahsilat"))
        self.assertNotIn("genel şartları", agent.tam_metin(m[1], CEVAP["sigortali"]["P-5"], "eposta", "", "T"))
        a = mesajlar("acente")
        self.assertNotIn("genel şartları", agent.tam_metin(a[0], CEVAP["acente"]["Örnek Acente A"], "sms", "", "T"))   # acenteye eklenmez


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_sigortali_eposta(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "h.xlsx"
            s = agent.calistir(YAS, cikti, KAD, BUGUN, odeme="Online ödeme", imza="Tahsilat Birimi", evet=True)
            self.assertEqual((len(MESAJLAR), s["tur"], s["kalan"]), (6, 1, []))
            for m in MESAJLAR:
                for ad in ("Kurgu Mobilya", "Hayali Lojistik", "Deneme Plastik"):
                    self.assertNotIn(ad, m)
            ws = load_workbook(cikti)["Mesajlar"]
            self.assertEqual([ws.cell(2, j).value for j in (2, 3, 4, 5, 6, 10)], ["Kurgu Mobilya Ltd. Şti.", "Son hatırlatma", 108, 90000, "P-2", "Geçti"])
            self.assertNotIn("durabilir", ws.cell(2, 8).value)
            self.assertIn("Konu: Son hatırlatma", s["txt"].read_text(encoding="utf-8"))

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_acente_sms_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(YAS, Path(tmp) / "h.xlsx", KAD, BUGUN, alici="acente", kanal="sms", evet=True)
            self.assertEqual((len(s["mesajlar"]), s["kalan"]), (3, []))
            self.assertTrue(all(len(agent.tam_metin(m, s["sonuc"][m["no"]], "sms", "", "Tahsilat Birimi")) <= agent.SMS_SINIR for m in s["mesajlar"]))
            with self.assertRaises(llm.LLMHatasi):                          # sabit kısımlar SMS'e sığmıyor
                agent.calistir(YAS, Path(tmp) / "h.xlsx", KAD, BUGUN, kanal="sms", odeme="x" * 300, evet=True)
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--tarih", "dün", "--evet"]), 2)
            self.assertEqual(agent.main(["--girdi", str(Path(tmp) / "yok.xlsx"), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
