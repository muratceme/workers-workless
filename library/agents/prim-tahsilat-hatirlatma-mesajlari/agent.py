"""
Prim Tahsilat Hatırlatma Mesajları — Workers / Workless AI Agent
Sigorta Tahsilat › Tahsilat Uzmanı (Acente Destek Uzmanı da kullanır)

Vadesi geçmiş prim taksitleri için sigortalıya (poliçe başına) ya da acenteye (acente başına) hatırlatma e-postası /
SMS taslakları hazırlar.
1. Kod açık taksit listesini okur ("Prim Alacak Yaşlandırma" kod bloğunun "Açık Taksitler" sayfası ya da aynı sütunlu bir
   tablo); vadesi geçmiş taksitleri poliçe ya da acente bazında toplar ve en uzun gecikmeye göre kademeyi (nazik / ikinci /
   son hatırlatma) belirler. Sigortalı adı sağlayıcıya gönderilmez.
2. Model kademenin tonuna uygun, kişiselleştirilmiş mesaj metnini yazar.
3. Kod mesajı denetler: her poliçe no, tutar ve vade tarihi kodun verdiği biçimde mesajda var mı; mesajda başka sayı var
   mı; hukuki sonuç ya da tehdit ifadesi (iptal, fesih, icra, dava, avukat, haciz, faiz…) var mı; SMS uzunluğu.
   Sorun varsa yalnız o mesaj modele geri gönderilir (en fazla --tur kez).
4. Kod hitabı, kademenin zorunlu cümlesini (kademeler dosyasında siz yazarsınız), ödeme bilgisini ve imzayı ekler.

Kullanım:
    python agent.py                                                      # örnek: sigortalılara e-posta
    python agent.py --alici acente --kanal sms
    python agent.py --girdi acik_taksitler.xlsx --kademeler kademeler.csv --odeme "IBAN: ... / online ödeme: ..." --imza "Tahsilat Birimi"
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import llm

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}
SMS_SINIR = 320
YASAK = re.compile(r"\b(iptal|fesih|feshed|teminat\w* (dur|son|kalk)|icra|dava|avukat|haciz|hukuki|yasal (takip|işlem|yollar)|kanuni|"
                   r"faiz|gecikme zammı|ceza|kara liste|mahkeme|noter|ihtar)\w*", re.I)


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1 or (s.count(".") == 1 and len(s.split(".")[1]) == 3):
        s = s.replace(".", "")
    try:
        return Decimal(s) if s not in ("", "-", ".") else None
    except InvalidOperation:
        return None


def tl(x: Decimal) -> str:
    return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(str(x or "").strip(), f).date()
        except ValueError:
            pass
    return None


def sayi_tokenlari(metin: str) -> set[str]:
    """Mesajdaki sayı ve tarih parçaları (karşılaştırma için kanonik metin)."""
    out = set()
    for t in re.findall(r"\d{1,2}\.\d{1,2}\.\d{4}", metin):
        out.add(t)
    temiz = re.sub(r"\d{1,2}\.\d{1,2}\.\d{4}", " ", metin)
    for x in re.findall(r"\d+(?:[.,]\d+)*", temiz):
        v = para(x)
        if v is not None:
            out.add(f"{v.normalize():f}")
    return out


# ----------------------------------------------------------------------------
# Girdi
# ----------------------------------------------------------------------------

SUTUN = {"police": ("police no", "police numarasi", "police"), "taksit": ("taksit no", "taksit"), "acente": ("acente", "acente adi"),
         "sigortali": ("sigortali", "sigortali adi", "musteri"), "vade": ("vade tarihi", "vade"), "kalan": ("kalan", "kalan tutar", "acik tutar"),
         "tutar": ("taksit tutari", "tutar"), "tahsil": ("tahsil edilen", "odenen"), "gecikme": ("gecikme gun", "gecikme (gun)", "gecikme"),
         "brans": ("brans",)}


def taksitleri_oku(yol: Path, bugun: date) -> list[dict]:
    """Açık taksitler: 'Açık Taksitler' sayfası varsa o, yoksa ilk sayfa. Gecikme sütunu yoksa vade ile bugünden hesaplanır."""
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: taksit dosyası bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, read_only=True, data_only=True)
        ws = wb["Açık Taksitler"] if "Açık Taksitler" in wb.sheetnames else wb.active
        s = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=metin.split("\n")[0].count)))
    s = [r for r in s if any(c not in (None, "") for c in r)]
    bas = next((i for i, r in enumerate(s[:10]) if "police no" in {katla(c) for c in r}), None)
    if bas is None:
        raise llm.LLMHatasi(f"{yol.name}: 'Poliçe No' başlıklı tablo bulunamadı.")
    b = [katla(c) for c in s[bas]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in SUTUN.items()}
    for a in ("police", "vade"):
        if idx[a] is None:
            raise llm.LLMHatasi(f"{yol.name}: '{SUTUN[a][0]}' sütunu yok.")
    if idx["kalan"] is None and idx["tutar"] is None:
        raise llm.LLMHatasi(f"{yol.name}: 'Kalan' ya da 'Taksit Tutarı' sütunu gerekli.")
    out = []
    for r in s[bas + 1:]:
        al = lambda a: None if idx[a] is None or idx[a] >= len(r) else r[idx[a]]  # noqa: E731
        v = tarih_coz(al("vade"))
        if not al("police") or v is None:
            continue
        kalan = para(al("kalan")) if idx["kalan"] is not None else (para(al("tutar")) or Decimal(0)) - (para(al("tahsil")) or Decimal(0))
        if not kalan or kalan <= 0:
            continue
        g = para(al("gecikme"))
        out.append({"police": str(al("police")).strip(), "taksit": str(al("taksit") or "").strip(), "acente": str(al("acente") or "").strip(),
                    "sigortali": str(al("sigortali") or "").strip(), "brans": str(al("brans") or "").strip(), "vade": v, "kalan": kalan,
                    "gecikme": int(g) if g is not None else (bugun - v).days})
    return out


def kademeleri_oku(yol: Path) -> list[dict]:
    metin = yol.read_text(encoding="utf-8-sig")
    s = list(csv.reader(metin.splitlines(), delimiter=";"))
    out = []
    for r in s[1:]:
        if len(r) < 3 or not r[0].strip():
            continue
        g = para(r[1])
        if g is None:
            raise llm.LLMHatasi(f"{yol.name}: '{r[0]}' kademesinin en az gecikme günü okunamadı.")
        out.append({"ad": r[0].strip(), "gun": int(g), "ton": r[2].strip(), "zorunlu": (r[3] if len(r) > 3 else "").strip()})
    if not out:
        raise llm.LLMHatasi(f"{yol.name}: kademe yok (Kademe;En Az Gecikme (gün);Ton;Zorunlu Cümle).")
    return sorted(out, key=lambda x: x["gun"])


def gruplar(taksitler: list[dict], alici: str, kademeler: list[dict]) -> list[dict]:
    """Vadesi geçmiş taksitleri poliçe (sigortalı) ya da acente bazında toplar; kademe en uzun gecikmeye göre."""
    g = defaultdict(list)
    for t in taksitler:
        if t["gecikme"] >= kademeler[0]["gun"]:
            g[t["police"] if alici == "sigortali" else (t["acente"] or "Acente belirtilmemiş")].append(t)
    out = []
    for i, (anahtar, ts) in enumerate(sorted(g.items(), key=lambda x: -max(t["gecikme"] for t in x[1])), 1):
        en = max(t["gecikme"] for t in ts)
        kad = [k for k in kademeler if en >= k["gun"]][-1]
        out.append({"no": f"M{i}", "alici_turu": alici, "anahtar": anahtar, "taksitler": sorted(ts, key=lambda t: (t["police"], t["vade"])), "kademe": kad, "en_uzun": en,
                    "toplam": sum((t["kalan"] for t in ts), Decimal(0)), "alici_ad": ts[0]["sigortali"] if alici == "sigortali" else anahtar})
    return out


def gerekli_degerler(m: dict) -> list[tuple[str, str]]:
    """Mesajda aynen geçmesi gereken değerler: poliçe no, toplam tutar, her vade tarihi."""
    out = [("toplam tutar", tl(m["toplam"]))]
    for p in dict.fromkeys(t["police"] for t in m["taksitler"]):
        out.append(("poliçe no", p))
    for v in dict.fromkeys(t["vade"] for t in m["taksitler"]):
        out.append(("vade", f"{v:%d.%m.%Y}"))
    return out


def veri_metni(m: dict, alici: str) -> str:
    l = [f"Alıcı türü: {'sigortalı' if alici == 'sigortali' else 'acente'}", f"Kademe: {m['kademe']['ad']} · Ton: {m['kademe']['ton']}",
         f"En uzun gecikme: {m['en_uzun']} gün · Vadesi geçmiş toplam: {tl(m['toplam'])} TL", "Vadesi geçmiş taksitler:"]
    for t in m["taksitler"]:
        parca = [f"Poliçe {t['police']}"]
        if t["brans"]:
            parca.append(t["brans"])
        if t["taksit"]:
            parca.append(f"taksit {t['taksit']}")
        parca += [f"vade {t['vade']:%d.%m.%Y}", f"kalan {tl(t['kalan'])} TL", f"{t['gecikme']} gün gecikme"]
        if alici == "acente":
            parca.append(f"sigortalı: [SİGORTALI-{_sig_no(m, t)}]")
        l.append("- " + " · ".join(parca))
    return "\n".join(l)


def _sig_no(m: dict, t: dict) -> int:
    return list(dict.fromkeys(x["sigortali"] for x in m["taksitler"])).index(t["sigortali"]) + 1


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

SEMA = {"type": "object", "properties": {"konu": S, "govde": S}, "required": ["konu", "govde"], "additionalProperties": False}


def denetle(m: dict, c: dict, kanal: str) -> list[str]:
    sorun, metin = [], c["konu"] + "\n" + c["govde"]
    for ad, v in gerekli_degerler(m):
        if v not in metin:
            sorun.append(f"{m['no']}: mesajda {ad} '{v}' yok (aynen bu biçimde yazın)")
    izinli = {f"{d.normalize():f}" for d in (m["toplam"], *(t["kalan"] for t in m["taksitler"]))} | {f"{t['vade']:%d.%m.%Y}" for t in m["taksitler"]}
    izinli |= {str(t["gecikme"]) for t in m["taksitler"]} | {str(m["en_uzun"])} | {x for t in m["taksitler"] for x in sayi_tokenlari(t["police"] + " " + t["taksit"])}
    fazla = sorted(sayi_tokenlari(metin) - izinli)
    if fazla:
        sorun.append(f"{m['no']}: veride olmayan sayı / tarih: {', '.join(fazla)}")
    if y := YASAK.search(metin):
        sorun.append(f"{m['no']}: hukuki sonuç / baskı ifadesi ('{y.group(0)}'); bu cümleleri yazmayın, gerekirse kod zorunlu cümleyi ekler")
    if re.match(r"\s*(sayın|değerli|merhaba|selam)", c["govde"], re.I):
        sorun.append(f"{m['no']}: gövde hitapla başlamasın; hitabı kod ekler")
    if re.search(r"\[SİGORTALI(-\d+)?\]", metin):
        sorun.append(f"{m['no']}: sigortalı yer tutucusu kullanılmış; sigortalıyı poliçe no ile anın")
    if kanal == "sms" and len(c["govde"]) > m["butce"]:
        sorun.append(f"{m['no']}: SMS gövdesi {len(c['govde'])} karakter; hitap, zorunlu cümle, ödeme bilgisi ve imzayla birlikte {m['sms_sinir']} "
                     f"karakteri geçmemesi için gövde en fazla {m['butce']} karakter olmalı")
    if kanal == "eposta" and not c["konu"].strip():
        sorun.append(f"{m['no']}: e-posta konusu boş")
    return sorun


def uret(mesajlar: list[dict], alici: str, kanal: str, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    sonuc = dict(onceki or {})
    hedef = mesajlar if duzeltme is None else [m for m in mesajlar if any(s.startswith(m["no"] + ":") for s in duzeltme)]
    for m in hedef:
        butce = f"<sms_butce>{m['butce']} karakter</sms_butce>\n" if kanal == "sms" else ""
        mesaj = f"<kanal>{'SMS' if kanal == 'sms' else 'e-posta'}</kanal>\n{butce}<veri>\n{veri_metni(m, alici)}\n</veri>"
        d = [s.split(": ", 1)[1] for s in (duzeltme or []) if s.startswith(m["no"] + ":")]
        if d:
            mesaj += ("\n<duzeltme>\nÖnceki mesaj: " + json.dumps(sonuc.get(m["no"]), ensure_ascii=False) + "\nKodun bulduğu sorunlar:\n"
                      + "\n".join(f"- {s}" for s in d) + "\n</duzeltme>")
        sonuc[m["no"]] = llm.json_iste(sistem, mesaj, SEMA)
    return sonuc


def duzeltme_dongusu(uret_f, denetle_f, tur: int) -> tuple[dict, list[str], int]:
    taslak = uret_f(None, None)
    kalan = denetle_f(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret_f(kalan, taslak)
        kalan = denetle_f(taslak)
    return taslak, kalan, yapilan


def tam_metin(m: dict, c: dict, kanal: str, odeme: str, imza: str) -> str:
    parca = [f"Sayın {m['alici_ad'] or 'Sigortalımız'},", c["govde"].strip()]
    if m["kademe"]["zorunlu"] and m["alici_turu"] == "sigortali":      # zorunlu cümle sigortalıya hitap eder
        parca.append(m["kademe"]["zorunlu"])
    if odeme:
        parca.append(f"Ödeme bilgisi: {odeme}")
    parca.append(f"Saygılarımızla,\n{imza}" if kanal == "eposta" else imza)
    return ("\n\n" if kanal == "eposta" else " ").join(parca)


# ----------------------------------------------------------------------------
# Çıktı
# ----------------------------------------------------------------------------

def rapor_yaz(cikti: Path, mesajlar: list[dict], sonuc: dict, kalan: list[str], tur: int, kanal: str, odeme: str, imza: str) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Mesajlar"
    bas = ["Mesaj", "Alıcı", "Kademe", "En Uzun Gecikme (gün)", "Vadesi Geçmiş Toplam", "Poliçeler", "Konu", "Mesaj Metni", "Karakter", "Son Kontrol"]
    ws.append(bas)
    for h in ws[1]:
        h.fill, h.font, h.alignment = BASLIK_DOLGU, BASLIK_YAZI, UST
    for j, w in enumerate((7, 30, 18, 10, 14, 24, 40, 90, 9, 30), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for m in mesajlar:
        c = sonuc[m["no"]]
        t = tam_metin(m, c, kanal, odeme, imza)
        sor = [s.split(": ", 1)[1] for s in kalan if s.startswith(m["no"] + ":")]
        ws.append([m["no"], m["alici_ad"], m["kademe"]["ad"], m["en_uzun"], float(m["toplam"]), ", ".join(dict.fromkeys(x["police"] for x in m["taksitler"])),
                   c["konu"] if kanal == "eposta" else "", t, len(t), "; ".join(sor) or "Geçti"])
        n = ws.max_row
        ws.cell(n, 5).number_format = "#,##0.00"
        ws.cell(n, 7).fill = ws.cell(n, 8).fill = AI_DOLGU
        ws.cell(n, 10).fill = KIRMIZI if sor else YESIL
        for h in ws[n]:
            h.alignment = UST
    k = wb.create_sheet("Kontrol")
    k.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    k.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    k.cell(2, 1).fill = KIRMIZI if kalan else YESIL
    k.append(["Önemli: Mesajlar taslaktır; göndermeden önce taksitlerin hâlâ açık olduğunu (son tahsilatlar) kontrol edin. Hukuki sonuç içeren "
              "cümleler yalnız kademeler dosyasındaki 'Zorunlu Cümle'den gelir; bu cümlenin poliçe genel şartlarınıza uygunluğunu siz onaylayın. "
              "Bu mesajlar resmî ihtar değildir."])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    txt = cikti.with_suffix(".txt")
    txt.write_text("\n\n".join(f"--- {m['no']} · {m['alici_ad']} · {m['kademe']['ad']} ---\n"
                               + (f"Konu: {sonuc[m['no']]['konu']}\n\n" if kanal == "eposta" else "") + tam_metin(m, sonuc[m["no"]], kanal, odeme, imza)
                               for m in mesajlar), encoding="utf-8")
    return txt


def calistir(girdi: Path, cikti: Path, kademe_yolu: Path, bugun: date, alici: str = "sigortali", kanal: str = "eposta", odeme: str = "",
             imza: str = "Tahsilat Birimi", tur: int = 2, evet: bool = False, sms_sinir: int = SMS_SINIR) -> dict:
    taksitler = taksitleri_oku(girdi, bugun)
    kademeler = kademeleri_oku(kademe_yolu)
    mesajlar = gruplar(taksitler, alici, kademeler)
    for m in mesajlar:
        m["sms_sinir"] = sms_sinir
        m["butce"] = sms_sinir - len(tam_metin(m, {"govde": ""}, "sms", odeme, imza))
        if kanal == "sms" and m["butce"] < 80:
            raise llm.LLMHatasi(f"{m['no']}: hitap, zorunlu cümle, ödeme bilgisi ve imza SMS'e sığmıyor ({sms_sinir - m['butce']} karakter). "
                                "Daha kısa --odeme / --imza verin ya da e-posta kullanın.")
    print(f"[OK] {len(taksitler)} açık taksit · {len(mesajlar)} mesaj ({'poliçe' if alici == 'sigortali' else 'acente'} başına)")
    sonuc, kalan, yapilan = {}, [], 0
    if mesajlar:
        llm.onay_al(f"{len(mesajlar)} mesaj için poliçe no, branş, vade, tutar ve gecikme günleri gönderilecek"
                    f"{'; acente adları' if alici == 'acente' else ''}. Sigortalı adları gönderilmez.", evet)
        sonuc, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(mesajlar, alici, kanal, d, o),
                                                 lambda x: [s for m in mesajlar for s in denetle(m, x[m["no"]], kanal)], tur)
    txt = rapor_yaz(cikti, mesajlar, sonuc, kalan, yapilan, kanal, odeme, imza)
    return {"mesajlar": mesajlar, "sonuc": sonuc, "kalan": kalan, "tur": yapilan, "txt": txt}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Vadesi geçmiş prim taksitleri için kademeli hatırlatma e-postası / SMS taslakları hazırlar.")
    p.add_argument("--girdi", type=Path, default=ORNEK / "prim_alacak_yaslandirma.xlsx", help="Açık taksitler (.xlsx/.csv) ya da Prim Alacak Yaşlandırma çıktısı")
    p.add_argument("--kademeler", type=Path, default=ORNEK / "kademeler.csv", help="Kademe;En Az Gecikme (gün);Ton;Zorunlu Cümle")
    p.add_argument("--alici", choices=["sigortali", "acente"], default="sigortali", help="Mesaj kime: poliçe başına sigortalıya ya da acente başına acenteye")
    p.add_argument("--kanal", choices=["eposta", "sms"], default="eposta")
    p.add_argument("--sms-sinir", type=int, default=SMS_SINIR, help=f"SMS mesajının toplam karakter sınırı (varsayılan {SMS_SINIR})")
    p.add_argument("--odeme", default="", help="Mesaja eklenecek ödeme bilgisi (IBAN, online ödeme bağlantısı vb.)")
    p.add_argument("--imza", default="Tahsilat Birimi")
    p.add_argument("--tarih", default=None, help="Gecikme sütunu yoksa hesap tarihi GG.AA.YYYY (varsayılan bugün; örnekte 01.10.2026)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "hatirlatma_mesajlari.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    bugun = tarih_coz(a.tarih) if a.tarih else (date(2026, 10, 1) if a.girdi == ORNEK / "prim_alacak_yaslandirma.xlsx" else date.today())
    if bugun is None:
        print(f"[X] --tarih '{a.tarih}' okunamadı (GG.AA.YYYY)")
        return 2
    try:
        s = calistir(a.girdi, a.cikti, a.kademeler, bugun, a.alici, a.kanal, a.odeme, a.imza, a.tur, a.evet, a.sms_sinir)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    for m in s["mesajlar"]:
        print(f"[OK] {m['no']} · {m['alici_ad']} · {m['kademe']['ad']} · {tl(m['toplam'])} TL · {m['en_uzun']} gün")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()} (+ {s['txt'].name})")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
