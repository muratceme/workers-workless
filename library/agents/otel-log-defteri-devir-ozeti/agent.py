"""
Otel Log Defteri Devir Özeti — Workers / Workless AI Agent
Otel Ön Büro › Night Auditor (Kat Şefi de kullanır)

Gece / vardiya log defteri kayıtlarını, gelen vardiyanın ve departmanların okuyacağı devir notuna çevirir: açık konular
(departman, oda, sorumlu, saat), kapananlar, misafir şikâyetleri, kasa / ödeme konuları.
1. Kod log dosyasını okur (ilk satır "Vardiya: … | Departman"; kayıtlar "SS:DD metin" biçiminde), satırları numaralar;
   telefon, e-posta, kimlik numarası ve --gizle ile verilen misafir adları maskelenir.
2. Model kayıtları maddelere ayırır: departman, oda, özet, durum (Açık / Kapandı / Bilgi), sorumlu, takip saati, tutar,
   öncelik; her madde dayandığı log satırlarını gösterir; kısa bir vardiya özeti yazar.
3. Kod denetler: her log satırı en az bir maddede kullanılmış mı; oda no, saat ve tutar dayanak satırlarda geçiyor mu;
   acil işaretli satırlar (VIP, transfer teyidi, yangın, su baskını, sağlık, güvenlik, kayıp eşya) "Acil" mi. Sorun varsa
   modele geri gönderilir (en fazla --tur kez).
4. Kod açık konuları departmana göre sıralar ve sorumlusu yazılmamış açık konuları "sahipsiz" diye işaretler.

Kullanım:
    python agent.py                                                    # örnek gece log defteri
    python agent.py --girdi log_15_10.txt --gizle "Ayşe Deneme"
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
RENK = {"Açık": "FFF4CE", "Kapandı": "E3F5E1", "Bilgi": "E8F0FE", "Acil": "FDE2E1"}
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}
DEPARTMAN = ["Ön Büro", "Kat Hizmetleri", "Teknik", "Yiyecek İçecek", "Güvenlik", "Muhasebe / Kasa", "Satış / Rezervasyon", "Diğer"]
ACIL = re.compile(r"\b(vip|transfer\w*|teyit etmedi|teyit edilmedi|yangin|duman|su baskini|su basti|ambulans|saglik|yarali|hastane|guvenlik ihlali|"
                  r"kavga|kayip esya|kayip cocuk|polis)\b")
SATIR = re.compile(r"^\s*(\d{1,2}[:.]\d{2})\s*[-–]?\s*(.+)$")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def para(x) -> Decimal | None:
    s = re.sub(r"[^\d,.\-]", "", str(x or ""))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return Decimal(s) if s not in ("", "-", ".") else None
    except InvalidOperation:
        return None


def saatler(metin: str) -> set[str]:
    return {f"{int(h):02d}:{m}" for h, m in re.findall(r"\b(\d{1,2})[:.](\d{2})\b", str(metin)) if int(h) < 24}


def sayilar(metin: str) -> set[Decimal]:
    temiz = re.sub(r"\b\d{1,2}[:.]\d{2}\b|\b\d{1,2}\.\d{1,2}\.\d{4}\b", " ", str(metin))
    return {v for x in re.findall(r"\d+(?:[.,]\d+)*", temiz) if (v := para(x)) is not None}


# ----------------------------------------------------------------------------
# Girdi
# ----------------------------------------------------------------------------

def logu_oku(yol: Path, kisiler: dict[str, str]) -> dict:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: log dosyası bulunamadı.")
    baslik, satirlar = "", []
    for s in belge.metin_oku(yol).splitlines():
        if not s.strip():
            continue
        if re.match(r"^\s*vardiya\s*:", s, re.I) and not baslik:
            baslik = s.split(":", 1)[1].strip()
            continue
        m = SATIR.match(s)
        if m:
            satirlar.append({"saat": f"{int(m[1][:-3]):02d}:{m[1][-2:]}", "metin": maskele(m[2].strip(), kisiler)})
        elif satirlar:
            satirlar[-1]["metin"] += " " + maskele(s.strip(), kisiler)
    if not satirlar:
        raise llm.LLMHatasi(f"{yol.name}: 'SS:DD metin' biçiminde log kaydı yok.")
    return {"baslik": maskele(baslik, kisiler), "satirlar": satirlar}


def takma_adlar(adlar: list[str]) -> dict[str, str]:
    return {a: f"[MİSAFİR-{i}]" for i, a in enumerate(dict.fromkeys(x.strip() for x in adlar if x.strip()), 1)}


def maskele(metin: str, kisiler: dict[str, str]) -> str:
    for g, t in sorted(kisiler.items(), key=lambda x: -len(x[0])):
        metin = re.sub(re.escape(g), t, metin, flags=re.I)
    return llm.maskele(metin)


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

MADDE = {"type": "object", "properties": {
    "departman": {"type": "string", "enum": DEPARTMAN}, "oda": S, "ozet": S, "durum": {"type": "string", "enum": ["Açık", "Kapandı", "Bilgi"]},
    "sorumlu": S, "takip_saati": S, "tutar": S, "oncelik": {"type": "string", "enum": ["Acil", "Normal"]},
    "misafir_sikayeti": {"type": "boolean"}, "satirlar": {"type": "array", "items": {"type": "integer"}}},
    "required": ["departman", "oda", "ozet", "durum", "sorumlu", "takip_saati", "tutar", "oncelik", "misafir_sikayeti", "satirlar"], "additionalProperties": False}
SEMA = {"type": "object", "properties": {"vardiya_ozeti": S, "maddeler": {"type": "array", "items": MADDE}}, "required": ["vardiya_ozeti", "maddeler"],
        "additionalProperties": False}


def denetle(log: dict, c: dict) -> list[str]:
    sorun, n = [], len(log["satirlar"])
    kullanilan = set()
    for i, m in enumerate(c["maddeler"], 1):
        on = f"Madde {i}"
        gecersiz = [k for k in m["satirlar"] if not 1 <= k <= n]
        if not m["satirlar"] or gecersiz:
            sorun.append(f"{on}: dayanak satırlar boş ya da geçersiz {gecersiz} (1-{n})")
            continue
        kullanilan |= set(m["satirlar"])
        kaynak = " ".join(f"{log['satirlar'][k - 1]['saat']} {log['satirlar'][k - 1]['metin']}" for k in m["satirlar"])
        if m["oda"].strip() and not set(re.findall(r"\d+", m["oda"])) <= set(re.findall(r"\d+", kaynak)):
            sorun.append(f"{on}: oda '{m['oda']}' dayanak satırlarda yok")
        if m["takip_saati"].strip() and not saatler(m["takip_saati"]) <= saatler(kaynak):
            sorun.append(f"{on}: takip saati '{m['takip_saati']}' dayanak satırlarda yazmıyor")
        if m["tutar"].strip() and (para(m["tutar"]) is None or para(m["tutar"]) not in sayilar(kaynak)):
            sorun.append(f"{on}: tutar '{m['tutar']}' dayanak satırlarda yazmıyor")
        fazla = sorted(sayilar(m["ozet"]) - sayilar(kaynak), key=float)
        if fazla:
            sorun.append(f"{on}: özette log satırlarında olmayan sayı: {', '.join(f'{x:g}' for x in fazla)}")
        if ACIL.search(katla(kaynak)) and m["durum"] == "Açık" and m["oncelik"] != "Acil":
            sorun.append(f"{on}: dayanak satırda acil işaret var ('{ACIL.search(katla(kaynak)).group(0)}'); açık konu 'Acil' olmalı")
    eksik = sorted(set(range(1, n + 1)) - kullanilan)
    if eksik:
        sorun.append(f"Log satırı {eksik} hiçbir maddede kullanılmamış")
    return sorun


def uret(log: dict, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    satir = "\n".join(f"{i}. {s['saat']} {s['metin']}" for i, s in enumerate(log["satirlar"], 1))
    mesaj = f"<vardiya>{log['baslik'] or 'belirtilmedi'}</vardiya>\n<log>\n{satir}\n</log>"
    if duzeltme:
        mesaj += ("\n<duzeltme>\nÖnceki özet: " + json.dumps(onceki, ensure_ascii=False) + "\nKodun bulduğu sorunlar:\n"
                  + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")
    return llm.json_iste(sistem, mesaj, SEMA)


def duzeltme_dongusu(uret_f, denetle_f, tur: int) -> tuple[dict, list[str], int]:
    taslak = uret_f(None, None)
    kalan = denetle_f(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret_f(kalan, taslak)
        kalan = denetle_f(taslak)
    return taslak, kalan, yapilan


def geri_ac(s, kisiler: dict[str, str]):
    if isinstance(s, str):
        for g, t in kisiler.items():
            s = s.replace(t, g)
        return s
    if isinstance(s, list):
        return [geri_ac(x, kisiler) for x in s]
    if isinstance(s, dict):
        return {k: geri_ac(v, kisiler) for k, v in s.items()}
    return s


# ----------------------------------------------------------------------------
# Çıktı
# ----------------------------------------------------------------------------

def devir_notu(log: dict, c: dict) -> str:
    acik = sorted([m for m in c["maddeler"] if m["durum"] == "Açık"], key=lambda m: (m["oncelik"] != "Acil", DEPARTMAN.index(m["departman"]), m["takip_saati"]))
    l = [f"DEVİR NOTU · {log['baslik']}", "", c["vardiya_ozeti"], "", f"AÇIK KONULAR ({len(acik)})"]
    for m in acik:
        sahip = m["sorumlu"] or "SORUMLU YOK — atanmalı"
        l.append(f"{'[ACİL] ' if m['oncelik'] == 'Acil' else ''}{m['departman']}{' · Oda ' + m['oda'] if m['oda'] else ''}: {m['ozet']}"
                 f" → {sahip}{' · ' + m['takip_saati'] if m['takip_saati'] else ''}{' · ' + m['tutar'] + ' TL' if m['tutar'] else ''}")
    for baslik, kosul in (("KAPANANLAR", lambda m: m["durum"] == "Kapandı"), ("BİLGİ", lambda m: m["durum"] == "Bilgi")):
        x = [m for m in c["maddeler"] if kosul(m)]
        if x:
            l += ["", baslik] + [f"{m['departman']}{' · Oda ' + m['oda'] if m['oda'] else ''}: {m['ozet']}" for m in x]
    sik = [m for m in c["maddeler"] if m["misafir_sikayeti"]]
    if sik:
        l += ["", f"MİSAFİR ŞİKÂYETLERİ ({len(sik)})"] + [f"Oda {m['oda'] or '-'}: {m['ozet']} ({m['durum']})" for m in sik]
    return "\n".join(l)


def rapor_yaz(cikti: Path, log: dict, c: dict, kalan: list[str], tur: int) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Devir"
    bas = ["Öncelik", "Durum", "Departman", "Oda", "Konu", "Sorumlu", "Takip Saati", "Tutar (TL)", "Şikâyet", "Log Satırları"]
    ws.append(bas)
    for h in ws[1]:
        h.fill, h.font, h.alignment = BASLIK_DOLGU, BASLIK_YAZI, UST
    for j, w in enumerate((8, 9, 16, 7, 60, 22, 9, 10, 8, 60), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    sira = sorted(c["maddeler"], key=lambda m: (["Açık", "Bilgi", "Kapandı"].index(m["durum"]), m["oncelik"] != "Acil", DEPARTMAN.index(m["departman"])))
    for m in sira:
        ws.append([m["oncelik"], m["durum"], m["departman"], m["oda"], m["ozet"], m["sorumlu"] or ("SORUMLU YOK" if m["durum"] == "Açık" else ""),
                   m["takip_saati"], None if para(m["tutar"]) is None else float(para(m["tutar"])), "Evet" if m["misafir_sikayeti"] else "",
                   "\n".join(f"{k}. {log['satirlar'][k - 1]['saat']} {log['satirlar'][k - 1]['metin']}" for k in m["satirlar"] if 1 <= k <= len(log["satirlar"]))])
        n = ws.max_row
        ws.cell(n, 2).fill = PatternFill("solid", fgColor=RENK[m["durum"]])
        if m["oncelik"] == "Acil" and m["durum"] == "Açık":
            ws.cell(n, 1).fill = PatternFill("solid", fgColor=RENK["Acil"])
        ws.cell(n, 3).fill = ws.cell(n, 5).fill = AI_DOLGU
        ws.cell(n, 6).fill = KIRMIZI if m["durum"] == "Açık" and not m["sorumlu"] else AI_DOLGU
        for h in ws[n]:
            h.alignment = UST
    k = wb.create_sheet("Kontrol")
    k.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    k.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti; bütün log satırları kullanıldı")])
    k.cell(2, 1).fill = KIRMIZI if kalan else YESIL
    k.append(["Önemli: Devir notu log defterinin yerine geçmez; log defteri asıl kayıttır. Sorumlusu olmayan açık konular gelen vardiya "
              "tarafından atanmalıdır."])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    txt = cikti.with_suffix(".txt")
    txt.write_text(devir_notu(log, c), encoding="utf-8")
    return txt


def calistir(girdi: Path, cikti: Path, gizle: list[str] | None = None, tur: int = 2, evet: bool = False) -> dict:
    kisiler = takma_adlar(gizle or [])
    log = logu_oku(girdi, kisiler)
    print(f"[OK] {log['baslik'] or girdi.name} · {len(log['satirlar'])} log kaydı")
    llm.onay_al(f"{len(log['satirlar'])} log kaydı gönderilecek (oda numaraları dahil; telefon, e-posta, TCKN maskeli"
                f"{f'; {len(kisiler)} misafir adı takma adla' if kisiler else ''}).", evet)
    c, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(log, d, o), lambda x: denetle(log, x), tur)
    c = geri_ac(c, kisiler)
    txt = rapor_yaz(cikti, log, c, kalan, yapilan)
    return {"log": log, "c": c, "kalan": kalan, "tur": yapilan, "txt": txt}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Otel log defteri kayıtlarını departman bazında devir notuna çevirir.")
    p.add_argument("--girdi", type=Path, default=ORNEK / "gece_log_defteri.txt", help="Log defteri (.txt/.docx): 'Vardiya: …' satırı ve 'SS:DD metin' kayıtları")
    p.add_argument("--gizle", nargs="*", default=[], metavar="AD", help="Gönderilmeyecek misafir / kişi adları")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "devir_notu.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.cikti, a.gizle, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    m = s["c"]["maddeler"]
    print(f"[OK] {len(m)} madde · açık {sum(x['durum'] == 'Açık' for x in m)} (acil {sum(x['durum'] == 'Açık' and x['oncelik'] == 'Acil' for x in m)}) · "
          f"sahipsiz {sum(x['durum'] == 'Açık' and not x['sorumlu'] for x in m)}")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()} (+ {s['txt'].name})")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
