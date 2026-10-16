"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

LOG = KLASOR / "ornek_veri" / "gece_log_defteri.txt"
CEVAP = json.loads((Path(__file__).parent / "sahte_cevaplar.json").read_text(encoding="utf-8"))
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    """İlk turda VIP transfer konusunu 'Normal' önceliğe koyar ve 06:10 halı kaydını atlar."""
    MESAJLAR.append(kullanici)
    return copy.deepcopy(CEVAP["dogru" if "<duzeltme>" in kullanici else "hatali"])


class KodTesti(unittest.TestCase):
    def test_okuma(self):
        log = agent.logu_oku(LOG, {})
        self.assertEqual(log["baslik"], "15.10.2026 23:00 - 16.10.2026 07:00 | Ön Büro - Gece")
        self.assertEqual([(s["saat"], s["metin"][:12]) for s in log["satirlar"]][:2], [("23:10", "512 numaralı"), ("23:45", "Grup Kurgu T")])
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "l.txt"
            y.write_text("Vardiya: Gece\n9.05 Misafir Ali Veli 210'a giriş yaptı\ndevamı: valiz hasarlı\n", encoding="utf-8")
            l = agent.logu_oku(y, agent.takma_adlar(["Ali Veli"]))
            self.assertEqual(l["satirlar"], [{"saat": "09:05", "metin": "Misafir [MİSAFİR-1] 210'a giriş yaptı devamı: valiz hasarlı"}])
            y.write_text("kayıt yok\n", encoding="utf-8")
            with self.assertRaises(llm.LLMHatasi):
                agent.logu_oku(y, {})

    def test_denetim(self):
        log = agent.logu_oku(LOG, {})
        self.assertEqual(agent.denetle(log, CEVAP["dogru"]), [])
        self.assertEqual(agent.denetle(log, CEVAP["hatali"]),
                         ["Madde 7: dayanak satırda acil işaret var ('vip'); açık konu 'Acil' olmalı", "Log satırı [8] hiçbir maddede kullanılmamış"])
        c = copy.deepcopy(CEVAP["dogru"])
        c["maddeler"][3].update(tutar="150")
        c["maddeler"][5].update(oda="619", takip_saati="10:00")
        c["maddeler"][0].update(ozet="Klima 2 saat çalışmadı", satirlar=[0])
        s = agent.denetle(log, c)
        for parca in ("Madde 4: tutar '150'", "Madde 6: oda '619'", "Madde 6: takip saati '10:00'", "Madde 1: dayanak satırlar boş ya da geçersiz [0]", "[1]"):
            self.assertTrue(any(parca in x for x in s), parca)

    def test_devir_notu(self):
        log = agent.logu_oku(LOG, {})
        t = agent.devir_notu(log, CEVAP["dogru"])
        acik = t.split("AÇIK KONULAR (6)\n")[1].split("\n\n")[0].splitlines()
        self.assertTrue(acik[0].startswith("[ACİL] Ön Büro · Oda 801") and acik[0].count("SORUMLU YOK") == 1)
        self.assertTrue(acik[1].startswith("Kat Hizmetleri:"))
        self.assertIn("MİSAFİR ŞİKÂYETLERİ (3)", t)


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ornek_ve_duzeltme(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "d.xlsx"
            self.assertEqual(agent.main(["--evet", "--cikti", str(cikti)]), 0)
            self.assertEqual(len(MESAJLAR), 2)
            self.assertIn("8. 06:10 Kat 3 koridor", MESAJLAR[0])
            ws = load_workbook(cikti)["Devir"]
            self.assertEqual([ws.cell(2, j).value for j in (1, 2, 4, 6, 7)], ["Acil", "Açık", "801", "SORUMLU YOK", "08:00"])
            self.assertEqual([ws.cell(r, 2).value for r in range(2, ws.max_row + 1)].count("Kapandı"), 2)
            self.assertIn("bütün log satırları kullanıldı", load_workbook(cikti)["Kontrol"].cell(2, 1).value)

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(agent.main(["--girdi", str(Path(tmp) / "yok.txt"), "--evet"]), 1)
            s = agent.calistir(LOG, Path(tmp) / "d.xlsx", tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 2)


if __name__ == "__main__":
    unittest.main()
