"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import re
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

FIHRIST, PROFIL, METINLER = agent.ORNEK / "fihrist_2026-10-09.txt", agent.ORNEK / "profil.txt", agent.ORNEK / "metinler"
MESAJLAR = []
SINIF = {"B001": "yuksek", "B002": "dusuk", "B003": "ilgisiz", "B004": "yuksek", "B005": "yuksek", "B006": "ilgisiz", "B007": "dusuk", "B008": "yuksek",
         "B009": "orta", "B010": "ilgisiz", "B011": "ilgisiz", "B012": "dusuk"}


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append((sistem, kullanici))
    if "<basliklar>" in kullanici:
        idler = re.findall(r'<baslik id="([^"]+)"', kullanici)
        return {"basliklar": [{"id": i, "ilgi": SINIF[i], "konu": "Konu " + i, "gerekce": "Gerekçe", "birimler": ["Çevre", "Uydurma Birim", "üretim"]}
                              for i in idler] + [{"id": "B999", "ilgi": "yuksek", "konu": "x", "gerekce": "x", "birimler": []}]}
    if "Ambalaj Atıklarının" in kullanici:
        return {"ozet": "Ambalaj bildirimi ve geri dönüştürülmüş hammadde zorunluluğu getiriliyor.",
                "degisiklikler": [
                    {"konu": "Yıllık bildirim", "aciklama": "5 inci maddeye bent eklendi.",
                     "alinti": "bir önceki yıl piyasaya sürdükleri ambalaj miktarını her yıl Mart ayı sonuna kadar"},
                    {"konu": "Geri dönüştürülmüş hammadde", "aciklama": "En az yüzde on.",
                     "alinti": "ürettikleri ambalajların en az yüzde onunda geri dönüştürülmüş hammadde kullanır"},
                    {"konu": "Uydurma ceza", "aciklama": "Ceza artırıldı.", "alinti": "idari para cezası iki katına çıkarılmıştır"}],
                "yapilacaklar": [{"is": "Bildirim sorumlusu belirlenmesi", "birim": "çevre", "zamanlama": "Her yıl Mart sonuna kadar"},
                                 {"is": "Hammadde planı", "birim": "Pazarlama", "zamanlama": "1/1/2027 öncesi"}],
                "belirsizlikler": ["Önceki hüküm metinde yok"]}
    return {"ozet": "Uzaktan eğitimin yarısı uygulamalı olacak.", "degisiklikler": [
        {"konu": "Uygulamalı eğitim", "aciklama": "11 inci madde.",
         "alinti": "Uzaktan eğitim yoluyla verilen eğitimlerin en az yarısı uygulamalı olarak işyerinde tamamlanır."}],
        "yapilacaklar": [{"is": "Eğitim planının güncellenmesi", "birim": "İSG", "zamanlama": "metinde süre yok"}], "belirsizlikler": []}


class KodTesti(unittest.TestCase):
    def test_fihrist(self):
        b = agent.fihrist_oku(FIHRIST)
        self.assertEqual(len(b), 12)                                                         # ilânlar hariç
        self.assertEqual((b[0].rg_tarih, b[0].rg_sayi, b[0].bolum, b[0].tur), (date(2026, 10, 9), "99001", "Yasama", "Kanun"))
        self.assertTrue(b[4].baslik.endswith("Değişiklik Yapılmasına Dair Yönetmelik (Kurgusal)"))   # iki satıra bölünmüş başlık
        self.assertEqual(b[1].tur, "Cumhurbaşkanı Kararları")
        self.assertEqual(len(agent.fihrist_oku(FIHRIST, ilanlar=True)), 14)

    def test_yururluk(self):
        rg = date(2026, 10, 9)
        y = agent.yururluk_tarihleri("MADDE 4 – (1) Bu Yönetmeliğin 3 üncü maddesi 1/1/2027 tarihinde, diğer hükümleri yayımı tarihinde yürürlüğe girer.", rg)
        self.assertEqual([t for t, _ in y], [date(2027, 1, 1), rg])       # açık tarih + "diğer hükümleri yayımı tarihinde"
        self.assertEqual(agent.yururluk_tarihleri("Bu Yönetmelik yayımını izleyen ayın başında yürürlüğe girer.", rg)[0][0], date(2026, 11, 1))
        self.assertEqual(agent.yururluk_tarihleri("Bu Tebliğ yayımı tarihinde yürürlüğe girer.", date(2026, 12, 20))[0][0], date(2026, 12, 20))
        self.assertIsNone(agent.yururluk_tarihleri("Bu Yönetmelik yayımından altı ay sonra yürürlüğe girer.", rg)[0][0])

    def test_profil_anahtar_tr(self):
        p = agent.profil_oku(PROFIL)
        self.assertIn("İSG", p["birimler"])
        b = agent.fihrist_oku(FIHRIST)
        agent.anahtar_tara(b, p["anahtar"])
        self.assertEqual(b[3].anahtar, ["ambalaj", "atık"])
        self.assertIn("enerji", b[10].anahtar)
        self.assertEqual(agent.tr_baslik("CUMHURBAŞKANI KARARI"), "Cumhurbaşkanı Kararı")
        self.assertTrue(agent.alinti_dogru("YÜZDE YETMİŞ ibaresi", 'Yönetmeliğin "yüzde yetmiş ibaresi" olarak'))

    def test_model_yok(self):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir([FIHRIST], PROFIL, Path(tmp) / "m.xlsx", metin_klasoru=METINLER, model=False, bugun=date(2026, 10, 9))
        self.assertEqual(MESAJLAR, [])
        self.assertEqual({b.id for b in agent.ilgili(s)}, {"B001", "B002", "B004", "B005", "B008", "B009", "B011"})
        turler = {(u["tur"], u["kim"]) for u in s["uyarilar"]}
        self.assertTrue({("Metin eşleşmedi", "ilgisiz_metin.txt"), ("Yakında yürürlük", "B004"), ("Yakında yürürlük", "B005")} <= turler)


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_uctan_uca(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "m.xlsx"
            s = agent.calistir([FIHRIST], PROFIL, cikti, metin_klasoru=METINLER, evet=True, bugun=date(2026, 10, 9))
            self.assertEqual(len(MESAJLAR), 3)                                    # 1 sınıflandırma + 2 özet (metni olan yüksek ilgililer)
            self.assertIn("## Başlık sınıflandırma", MESAJLAR[0][0])
            self.assertNotIn("## Düzenleme özeti", MESAJLAR[0][0])
            self.assertIn("<yururluk_maddesi>", MESAJLAR[1][1])
            b = {x.id: x for x in s["basliklar"]}
            self.assertEqual(b["B004"].sinif["birimler"], ["Çevre", "Üretim"])      # uydurma birim atıldı, yazım düzeldi
            self.assertEqual([d["konu"] for d in b["B004"].ozet["degisiklikler"]], ["Yıllık bildirim", "Geri dönüştürülmüş hammadde"])
            self.assertEqual([x["birim"] for x in b["B004"].ozet["yapilacaklar"]], ["Çevre", "Belirlenecek"])
            self.assertEqual([t for t, _ in b["B004"].yururluk], [date(2027, 1, 1), date(2026, 10, 9)])
            self.assertEqual(b["B005"].yururluk[0][0], date(2026, 11, 1))
            self.assertNotIn("B999", b)
            turler = {(u["tur"], u["kim"]) for u in s["uyarilar"]}
            self.assertTrue({("Alıntı doğrulanamadı", "B004"), ("Tam metin gerekli", "B001"), ("Tam metin gerekli", "B008"),
                             ("Anahtar kelime / ilgisiz", "B011"), ("Yakında yürürlük", "B005")} <= turler)
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Özet", "İlgili Düzenlemeler", "Değişiklikler", "Yapılacaklar", "Yürürlük Takvimi", "Tüm Başlıklar", "Uyarılar"])
            self.assertEqual(wb["Değişiklikler"].max_row, 4)
            bulten = s["bulten"].read_text(encoding="utf-8")
            self.assertIn("yüzde onunda geri dönüştürülmüş", bulten)
            self.assertNotIn("iki katına", bulten)

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(agent.main(["--model-yok", "--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(agent.main(["--model-yok", "--bugun", "x"]), 2)
            self.assertEqual(agent.main(["--fihrist", str(Path(tmp) / "yok.txt")]), 1)


if __name__ == "__main__":
    unittest.main()
