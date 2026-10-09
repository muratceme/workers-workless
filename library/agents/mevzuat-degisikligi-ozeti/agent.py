"""
Mevzuat Değişikliği Özeti — Workers / Workless AI Agent
Hukuk › Hukuk Müdürü

Resmî Gazete fihristini şirketin faaliyet alanına göre süzer ve ilgili düzenlemelerin etkisini özetler:
1. Kod fihristi (Resmî Gazete'nin günlük içindekiler sayfası, metin olarak kaydedilmiş) okur: tarih, sayı, bölüm
   (Yasama / Yürütme ve İdare / Yargı / İlân), tür (Kanun, Cumhurbaşkanı Kararı, Yönetmelik, Tebliğ...) ve başlık.
   İlân bölümü varsayılan olarak dışarıda bırakılır.
2. Model başlıkları şirket profiline göre sınıflandırır (yüksek / orta / düşük / ilgisiz), konu, gerekçe ve etkilenen
   birimleri yazar. Kod kimlikleri ve birim adlarını doğrular; profildeki anahtar kelimelerden biri başlıkta geçtiği
   hâlde "ilgisiz" denen başlığı kontrol listesine ekler.
3. Tam metni verilen ilgili düzenlemeler için kod yürürlük maddesini metinden birebir çıkarır ve yürürlük tarihlerini
   hesaplar ("yayımı tarihinde", "1/1/2027 tarihinde", "yayımını izleyen ayın başında"). Model özeti, değişiklikleri ve
   yapılacakları yazar; her değişiklik için metinden birebir alıntı verir. Kod alıntının metinde geçtiğini doğrular;
   doğrulanamayan değişiklik rapora alınmaz.
Model kullanılmadan (--model-yok) yalnız fihrist ve anahtar kelime taraması yapılır.
İnternetten metin indirmez: fihrist ve tam metinler kullanıcı tarafından kaydedilir.

Kullanım:
    python agent.py                                                     # örnek: kurgusal fihrist + 2 tam metin
    python agent.py --fihrist rg_2026-10-09.txt --metinler metinler/ --profil profil.txt
    python agent.py --fihrist rg_*.txt --profil profil.txt --model-yok
"""
from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import llm

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
PAKET = 60
AZAMI_METIN = 120_000
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
ILGI = ["yuksek", "orta", "dusuk", "ilgisiz"]
ILGI_ADI = {"yuksek": "Yüksek", "orta": "Orta", "dusuk": "Düşük", "ilgisiz": "İlgisiz", "anahtar": "Anahtar kelime"}
BOLUMLER = {"yasama bolumu": "Yasama", "yurutme ve idare bolumu": "Yürütme ve İdare", "yargi bolumu": "Yargı", "ilan bolumu": "İlân"}
MADDE_BASI = re.compile(r"(?m)^\s*(?:MADDE|Madde|GEÇİCİ MADDE|Geçici Madde|EK MADDE)\s+\d+")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def _normal(s: str) -> str:
    return " ".join(str(s).replace("İ", "i").replace("I", "ı").replace("’", "'").replace("‘", "'").replace("“", '"').replace("”", '"')
                    .lower().split()).strip(" .,;:!?\"'")


def alinti_dogru(alinti: str, kaynak: str) -> bool:
    a = _normal(alinti)
    return len(a) >= 12 and a in _normal(kaynak)


def metin_oku(yol: Path) -> str:
    for kod in ("utf-8-sig", "cp1254"):
        try:
            return yol.read_text(encoding=kod)
        except UnicodeDecodeError:
            continue
    raise ValueError(f"{yol.name}: metin okunamadı (UTF-8 veya Windows-1254 bekleniyor)")


# ----------------------------------------------------------------------------
# Profil ve fihrist
# ----------------------------------------------------------------------------

def profil_oku(yol: Path) -> dict:
    """'Alan: değer' satırları. Birimler ve Anahtar kelimeler ';' veya ',' ile ayrılır."""
    p = {"satirlar": [], "birimler": [], "anahtar": []}
    for satir in metin_oku(yol).splitlines():
        if not satir.strip():
            continue
        p["satirlar"].append(satir.strip())
        ad, _, deger = satir.partition(":")
        k = katla(ad)
        if k.startswith("birim"):
            p["birimler"] = [x.strip() for x in re.split(r"[;,]", deger) if x.strip()]
        elif k.startswith("anahtar"):
            p["anahtar"] = [x.strip() for x in re.split(r"[;,]", deger) if x.strip()]
    if not p["birimler"]:
        raise ValueError(f"{yol.name}: 'Birimler:' satırı bulunamadı")
    return p


@dataclass
class Baslik:
    id: str
    rg_tarih: date | None
    rg_sayi: str
    bolum: str
    tur: str
    baslik: str
    anahtar: list[str] = field(default_factory=list)
    sinif: dict = field(default_factory=dict)
    metin: str = ""
    metin_dosyasi: str = ""
    yururluk_maddesi: str = ""
    yururluk: list[tuple[date | None, str]] = field(default_factory=list)
    ozet: dict = field(default_factory=dict)
    kontrol: list[str] = field(default_factory=list)

    @property
    def ilgi(self) -> str:
        return self.sinif.get("ilgi", "anahtar" if self.anahtar else "ilgisiz")


def tr_baslik(s: str) -> str:
    """'CUMHURBAŞKANI KARARI' → 'Cumhurbaşkanı Kararı' (Türkçe i/ı doğru)."""
    kucuk = s.replace("İ", "i").replace("I", "ı").lower()
    return " ".join(k[:1].replace("i", "İ").replace("ı", "I").upper() + k[1:] for k in kucuk.split())


def _baslik_mi(satir: str) -> bool:
    harf = re.sub(r"[^A-Za-zÇĞİÖŞÜÂÎÛçğıöşüâîû]", "", satir)
    return len(harf) >= 4 and harf == harf.upper() and not re.match(r"^\s*[—–-]", satir)


def fihrist_oku(yol: Path, ilanlar: bool = False, baslangic_no: int = 1) -> list[Baslik]:
    metin = metin_oku(yol)
    m = re.search(r"(\d{1,2})[./](\d{1,2})[./](\d{4})", metin[:500]) or re.search(r"(\d{4})-(\d{2})-(\d{2})", yol.stem)
    rg_tarih = None
    if m:
        a, b, c = m.groups()
        rg_tarih = date(int(a), int(b), int(c)) if len(a) == 4 else date(int(c), int(b), int(a))
    sm = re.search(r"Sayı\s*:?\s*(\d+)", metin[:500], re.I)
    bolum, tur, liste, no = "", "", [], baslangic_no
    devam = False
    for ham in metin.splitlines():
        satir = ham.strip()
        if not satir:
            continue
        k = katla(satir)
        if k in BOLUMLER:
            bolum, tur, devam = BOLUMLER[k], "", False
            continue
        if re.match(r"^[—–-]+\s*", satir):
            liste.append([bolum, tur, re.sub(r"^[—–-]+\s*", "", satir)])
            devam = True
            continue
        if re.match(r"^[a-zçğıöşü]\s*[-–)]\s", satir) or (_baslik_mi(satir) and not k.startswith("resmi gazete")):
            tur, devam = tr_baslik(re.sub(r"^[a-zçğıöşü]\s*[-–)]\s*", "", satir)), False     # alt başlık (ör. "a - Yargı İlânları")
            continue
        if devam:
            liste[-1][2] += " " + satir        # birden çok satıra bölünmüş başlık
    sonuc = []
    for b, t, bas in liste:
        if b == "İlân" and not ilanlar:
            continue
        sonuc.append(Baslik(f"B{no:03d}", rg_tarih, sm.group(1) if sm else "", b or "—", t or "—", " ".join(bas.split())))
        no += 1
    return sonuc


def anahtar_tara(basliklar: list[Baslik], anahtarlar: list[str]) -> None:
    for b in basliklar:
        kb = " " + katla(b.baslik) + " "
        b.anahtar = [a for a in anahtarlar if katla(a) and " " + katla(a) in kb]


def metinleri_esle(basliklar: list[Baslik], klasor: Path | None) -> list[dict]:
    uy = []
    if not klasor:
        return uy
    for yol in sorted(klasor.glob("*.txt")):
        t = metin_oku(yol)
        ilk = next((s.strip() for s in t.splitlines() if s.strip()), "")
        a = set(katla(ilk).split())
        en, puan = None, 0.0
        for b in basliklar:
            x = set(katla(b.baslik).split())
            j = len(a & x) / len(a | x) if a | x else 0
            if j > puan:
                en, puan = b, j
        if en is None or puan < 0.6:
            uy.append({"onem": "Orta", "tur": "Metin eşleşmedi", "kim": yol.name, "aciklama": f"İlk satır fihristteki hiçbir başlıkla eşleşmedi: '{ilk[:80]}'"})
            continue
        if len(t) > AZAMI_METIN:
            uy.append({"onem": "Orta", "tur": "Metin kısaltıldı", "kim": yol.name, "aciklama": f"{len(t):,} karakter; ilk {AZAMI_METIN:,} karakter kullanıldı"})
            t = t[:AZAMI_METIN]
        en.metin, en.metin_dosyasi = t, yol.name
    return uy


# ----------------------------------------------------------------------------
# Yürürlük (kod)
# ----------------------------------------------------------------------------

def maddeler(metin: str) -> list[str]:
    bas = [m.start() for m in MADDE_BASI.finditer(metin)]
    return [metin[a:b].strip() for a, b in zip(bas, bas[1:] + [len(metin)])]


def yururluk_maddesi(metin: str) -> str:
    adaylar = [m for m in maddeler(metin) if "yürürlüğe gir" in m.lower() or "yururluge gir" in katla(m)]
    return "\n".join(adaylar)


def yururluk_tarihleri(madde: str, rg_tarih: date | None) -> list[tuple[date | None, str]]:
    sonuc = []
    for cumle in re.split(r"(?<=[.;])\s+|\n", madde):
        c = cumle.strip()
        k = katla(c)
        if "yururluge gir" not in k:
            continue
        tarihler = [date(int(y), int(m), int(d)) for d, m, y in re.findall(r"\b(\d{1,2})[./](\d{1,2})[./](\d{4})\b", c)]
        if tarihler:
            sonuc += [(t, c) for t in tarihler]
            if ("yayimi tarihinde" in k or "yayimlandigi tarihte" in k) and rg_tarih:
                sonuc.append((rg_tarih, c))          # "3 üncü maddesi 1/1/2027 tarihinde, diğer hükümleri yayımı tarihinde"
        elif "yayimini izleyen ayin basinda" in k and rg_tarih:
            sonuc.append((date(rg_tarih.year + rg_tarih.month // 12, rg_tarih.month % 12 + 1, 1), c))
        elif ("yayimi tarihinde" in k or "yayimlandigi tarihte" in k) and rg_tarih:
            sonuc.append((rg_tarih, c))
        else:
            sonuc.append((None, c))
    return sonuc


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

SEMA_SINIF = {
    "type": "object",
    "properties": {"basliklar": {"type": "array", "items": {
        "type": "object",
        "properties": {"id": {"type": "string"}, "ilgi": {"type": "string", "enum": ILGI}, "konu": {"type": "string"}, "gerekce": {"type": "string"},
                       "birimler": {"type": "array", "items": {"type": "string"}}},
        "required": ["id", "ilgi", "konu", "gerekce", "birimler"], "additionalProperties": False}}},
    "required": ["basliklar"], "additionalProperties": False,
}
SEMA_OZET = {
    "type": "object",
    "properties": {
        "ozet": {"type": "string"},
        "degisiklikler": {"type": "array", "items": {"type": "object", "properties": {"konu": {"type": "string"}, "aciklama": {"type": "string"},
                                                                                      "alinti": {"type": "string"}},
                                                     "required": ["konu", "aciklama", "alinti"], "additionalProperties": False}},
        "yapilacaklar": {"type": "array", "items": {"type": "object", "properties": {"is": {"type": "string"}, "birim": {"type": "string"},
                                                                                    "zamanlama": {"type": "string"}},
                                                    "required": ["is", "birim", "zamanlama"], "additionalProperties": False}},
        "belirsizlikler": {"type": "array", "items": {"type": "string"}},
    },
    "required": ["ozet", "degisiklikler", "yapilacaklar", "belirsizlikler"], "additionalProperties": False,
}


def istem(bolum: str) -> str:
    """prompt.md: ortak giriş + '## Başlık sınıflandırma' veya '## Düzenleme özeti' bölümü."""
    t = (BURASI / "prompt.md").read_text(encoding="utf-8")
    giris, _, kalan = t.partition("## Başlık sınıflandırma")
    sinif, _, ozet = kalan.partition("## Düzenleme özeti")
    return giris + ("## Başlık sınıflandırma" + sinif if bolum == "sinif" else "## Düzenleme özeti" + ozet)


def _profil_xml(p: dict) -> str:
    return "<sirket_profili>\n" + "\n".join(p["satirlar"]) + "\n</sirket_profili>"


def birim_dogrula(adlar: list[str], birimler: list[str]) -> list[str]:
    s = {katla(b): b for b in birimler}
    return list(dict.fromkeys(s[katla(a)] for a in adlar if katla(a) in s))


def siniflandir(basliklar: list[Baslik], profil: dict) -> list[dict]:
    uy = []
    sistem = istem("sinif")
    for i in range(0, len(basliklar), PAKET):
        parca = basliklar[i:i + PAKET]
        mesaj = "\n".join([_profil_xml(profil), "<birim_listesi>" + "; ".join(profil["birimler"]) + "</birim_listesi>", "<basliklar>",
                           *[f'<baslik id="{b.id}" bolum="{b.bolum}" tur="{b.tur}">{b.baslik}</baslik>' for b in parca], "</basliklar>"])
        yanit = llm.json_iste(sistem, mesaj, SEMA_SINIF)
        harita = {b.id: b for b in parca}
        for s in yanit.get("basliklar", []):
            b = harita.get(s.get("id"))
            if b is None or b.sinif:
                continue
            s["birimler"] = birim_dogrula(s.get("birimler", []), profil["birimler"])
            b.sinif = s
    for b in basliklar:
        if not b.sinif:
            b.kontrol.append("Model bu başlık için sonuç döndürmedi")
            uy.append({"onem": "Orta", "tur": "Sınıflandırılmadı", "kim": b.id, "aciklama": f"{b.baslik[:90]} — elle değerlendirin"})
        elif b.sinif["ilgi"] == "ilgisiz" and b.anahtar:
            b.kontrol.append(f"Model ilgisiz dedi ama anahtar kelime geçiyor: {', '.join(b.anahtar)}")
            uy.append({"onem": "Bilgi", "tur": "Anahtar kelime / ilgisiz", "kim": b.id,
                       "aciklama": f"{b.baslik[:90]} — başlıkta '{', '.join(b.anahtar)}' geçiyor; kontrol edin"})
    return uy


def ozetle(b: Baslik, profil: dict) -> list[dict]:
    uy = []
    mesaj = "\n".join([_profil_xml(profil), "<birim_listesi>" + "; ".join(profil["birimler"]) + "</birim_listesi>",
                       f'<duzenleme baslik="{b.baslik}" tur="{b.tur}" rg_tarihi="{b.rg_tarih:%d.%m.%Y}">' if b.rg_tarih else f'<duzenleme baslik="{b.baslik}">',
                       b.metin, "</duzenleme>",
                       "<yururluk_maddesi>\n" + (b.yururluk_maddesi or "(bulunamadı)") + "\n</yururluk_maddesi>"])
    y = llm.json_iste(istem("ozet"), mesaj, SEMA_OZET, max_tokens=12000)
    dogru = []
    for d in y.get("degisiklikler", []):
        if alinti_dogru(d.get("alinti", ""), b.metin):
            dogru.append(d)
        else:
            b.kontrol.append(f"Değişiklik çıkarıldı (alıntı metinde yok): {d.get('konu', '')}")
            uy.append({"onem": "Bilgi", "tur": "Alıntı doğrulanamadı", "kim": b.id, "aciklama": f"'{d.get('konu', '')}' maddesi rapora alınmadı"})
    for x in y.get("yapilacaklar", []):
        bir = birim_dogrula([x.get("birim", "")], profil["birimler"])
        x["birim"] = bir[0] if bir else "Belirlenecek"
    y["degisiklikler"] = dogru
    b.ozet = y
    return uy


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
ILGI_RENK = {"yuksek": "FDE2E1", "orta": "FFF4CE", "dusuk": "E8F0FE", "anahtar": "FFF4CE"}
AI = PatternFill("solid", fgColor="EDE7FF")
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, basliklar, genislik=()):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def ilgili(s: dict) -> list[Baslik]:
    sira = {"yuksek": 0, "orta": 1, "anahtar": 2}
    return sorted([b for b in s["basliklar"] if b.ilgi in sira], key=lambda b: (sira[b.ilgi], b.id))


def rapor_yaz(cikti: Path, s: dict) -> None:
    wb = Workbook()
    oz = wb.active
    oz.title = "Özet"
    _baslik(oz, ["Gösterge", "Değer"], (44, 50))
    tarihler = sorted({b.rg_tarih for b in s["basliklar"] if b.rg_tarih})
    for k, v in [("Resmî Gazete tarihi", ", ".join(f"{t:%d.%m.%Y}" for t in tarihler) or "—"), ("Başlık sayısı (ilânlar hariç)", len(s["basliklar"])),
                 ("Yüksek ilgili", sum(b.ilgi == "yuksek" for b in s["basliklar"])), ("Orta ilgili", sum(b.ilgi == "orta" for b in s["basliklar"])),
                 ("Anahtar kelime eşleşmesi", sum(bool(b.anahtar) for b in s["basliklar"])),
                 ("Tam metni özetlenen", sum(bool(b.ozet) for b in s["basliklar"])),
                 ("İlgili ama tam metni olmayan", sum(not b.metin for b in ilgili(s))), ("Model", s["model"])]:
        oz.append([k, v])
    oz.append([])
    oz.append(["Mor hücreler yapay zekâ tarafından yazılmıştır; hukuk biriminin onayından önce bağlayıcı değildir."])

    il = wb.create_sheet("İlgili Düzenlemeler")
    _baslik(il, ["No", "RG Tarihi", "Tür", "Başlık", "İlgi", "Konu", "Gerekçe", "Birimler", "Özet", "Yürürlük", "Tam Metin", "Karar / Sorumlu"],
            (6, 11, 16, 50, 9, 22, 40, 22, 70, 30, 14, 20))
    for b in ilgili(s):
        yur = "; ".join(f"{t:%d.%m.%Y}" if t else "metne bakın" for t, _ in b.yururluk)
        il.append([b.id, b.rg_tarih, b.tur, b.baslik, ILGI_ADI[b.ilgi], b.sinif.get("konu", ""), b.sinif.get("gerekce", ""), ", ".join(b.sinif.get("birimler", [])),
                   b.ozet.get("ozet", "" if b.metin else "Tam metin verilmedi"), yur, b.metin_dosyasi or "yok", ""])
        il.cell(il.max_row, 2).number_format = "DD.MM.YYYY"
        il.cell(il.max_row, 5).fill = PatternFill("solid", fgColor=ILGI_RENK[b.ilgi])
        for j in (6, 7, 8, 9):
            il.cell(il.max_row, j).alignment = UST
            if b.sinif or j == 9 and b.ozet:
                il.cell(il.max_row, j).fill = AI
        il.cell(il.max_row, 4).alignment = UST
        il.cell(il.max_row, 12).fill = PatternFill("solid", fgColor="FFF4CE")

    de = wb.create_sheet("Değişiklikler")
    _baslik(de, ["No", "Düzenleme", "Konu", "Açıklama", "Metinden Alıntı (doğrulandı)"], (6, 40, 24, 60, 70))
    for b in ilgili(s):
        for d in b.ozet.get("degisiklikler", []):
            de.append([b.id, b.baslik, d["konu"], d["aciklama"], d["alinti"]])
            for j in (2, 3, 4, 5):
                de.cell(de.max_row, j).alignment = UST
            for j in (3, 4):
                de.cell(de.max_row, j).fill = AI

    ya = wb.create_sheet("Yapılacaklar")
    _baslik(ya, ["No", "Düzenleme", "Yapılacak", "Birim", "Zamanlama", "Sorumlu", "Durum"], (6, 40, 60, 18, 24, 16, 12))
    for b in ilgili(s):
        for x in b.ozet.get("yapilacaklar", []):
            ya.append([b.id, b.baslik, x["is"], x["birim"], x["zamanlama"], "", ""])
            for j in (2, 3, 5):
                ya.cell(ya.max_row, j).alignment = UST
            for j in (3, 4, 5):
                ya.cell(ya.max_row, j).fill = AI
            for j in (6, 7):
                ya.cell(ya.max_row, j).fill = PatternFill("solid", fgColor="FFF4CE")

    yt = wb.create_sheet("Yürürlük Takvimi")
    _baslik(yt, ["Yürürlük Tarihi", "Kalan Gün", "No", "Düzenleme", "Yürürlük Hükmü (metinden)"], (12, 9, 6, 50, 90))
    satirlar = [(t, b, c) for b in s["basliklar"] for t, c in b.yururluk]
    for t, b, c in sorted(satirlar, key=lambda x: (x[0] or date.max, x[1].id)):
        yt.append([t, (t - s["bugun"]).days if t else None, b.id, b.baslik, c])
        yt.cell(yt.max_row, 1).number_format = "DD.MM.YYYY"
        yt.cell(yt.max_row, 4).alignment = yt.cell(yt.max_row, 5).alignment = UST

    tb = wb.create_sheet("Tüm Başlıklar")
    _baslik(tb, ["No", "RG Tarihi", "Sayı", "Bölüm", "Tür", "Başlık", "Anahtar Kelime", "İlgi", "Konu", "Kontrol"], (6, 11, 8, 16, 18, 70, 18, 10, 24, 40))
    for b in s["basliklar"]:
        tb.append([b.id, b.rg_tarih, b.rg_sayi, b.bolum, b.tur, b.baslik, ", ".join(b.anahtar), ILGI_ADI[b.ilgi], b.sinif.get("konu", ""), "; ".join(b.kontrol)])
        tb.cell(tb.max_row, 2).number_format = "DD.MM.YYYY"
        tb.cell(tb.max_row, 6).alignment = UST
        if b.ilgi in ILGI_RENK:
            tb.cell(tb.max_row, 8).fill = PatternFill("solid", fgColor=ILGI_RENK[b.ilgi])
    tb.auto_filter.ref = f"A1:J{tb.max_row}"

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "No / Dosya", "Açıklama"], (9, 24, 18, 100))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["kim"], u["aciklama"]])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def bulten_yaz(yol: Path, s: dict) -> None:
    tarihler = sorted({b.rg_tarih for b in s["basliklar"] if b.rg_tarih})
    t = [f"# Mevzuat Bülteni — Resmî Gazete {', '.join(f'{x:%d.%m.%Y}' for x in tarihler)}", "",
         "> Bu bülten yapay zekâ desteğiyle hazırlanmış bir taslaktır. Alıntılar metinden birebir doğrulanmıştır; değerlendirmeler",
         "> hukuk biriminin onayından önce bağlayıcı değildir.", ""]
    lst = ilgili(s)
    if not lst:
        t.append("Şirketi ilgilendiren düzenleme bulunmadı.")
    for b in lst:
        t += [f"## {b.baslik}", "", f"- **Tür:** {b.tur} · **İlgi:** {ILGI_ADI[b.ilgi]}" + (f" · **Birimler:** {', '.join(b.sinif.get('birimler', []))}" if b.sinif else "")]
        if b.sinif.get("gerekce"):
            t.append(f"- **Neden önemli:** {b.sinif['gerekce']}")
        hukumler = {}
        for tar, c in b.yururluk:
            hukumler.setdefault(c, []).append(tar)
        for c, tarihler in hukumler.items():
            yazi = ", ".join(f"{x:%d.%m.%Y}" for x in tarihler if x)
            t.append(f"- **Yürürlük:** {yazi} — „{c}”" if yazi else f"- **Yürürlük:** „{c}”")
        t.append("")
        if b.ozet:
            t += [b.ozet["ozet"], ""]
            if b.ozet["degisiklikler"]:
                t.append("**Değişiklikler**")
                t += [f"- {d['konu']}: {d['aciklama']} — „{d['alinti']}”" for d in b.ozet["degisiklikler"]]
                t.append("")
            if b.ozet["yapilacaklar"]:
                t.append("**Yapılacaklar**")
                t += [f"- [{x['birim']}] {x['is']} ({x['zamanlama']})" for x in b.ozet["yapilacaklar"]]
                t.append("")
        elif not b.metin:
            t += ["_Tam metin verilmedi; Resmî Gazete'den metni kaydedip yeniden çalıştırın._", ""]
    yol.write_text("\n".join(t) + "\n", encoding="utf-8")


def calistir(fihristler: list[Path], profil_yolu: Path, cikti: Path, *, metin_klasoru: Path | None = None, model: bool = True, ilanlar: bool = False,
             evet: bool = False, bugun: date | None = None) -> dict:
    profil = profil_oku(profil_yolu)
    basliklar = []
    for f in fihristler:
        basliklar += fihrist_oku(f, ilanlar, len(basliklar) + 1)
    if not basliklar:
        raise llm.LLMHatasi("Fihristte başlık bulunamadı. Başlıklar '—' ile başlayan satırlar olmalıdır.")
    uy = [{"onem": "Orta", "tur": "Fihrist tarihi yok", "kim": f.name, "aciklama": "Resmî Gazete tarihi okunamadı; yürürlük tarihleri hesaplanamaz"}
          for f in fihristler if not any(b.rg_tarih for b in fihrist_oku(f, True))]
    anahtar_tara(basliklar, profil["anahtar"])
    uy += metinleri_esle(basliklar, metin_klasoru)
    for b in basliklar:
        if b.metin:
            b.yururluk_maddesi = yururluk_maddesi(b.metin)
            b.yururluk = yururluk_tarihleri(b.yururluk_maddesi, b.rg_tarih)
            if not b.yururluk_maddesi:
                uy.append({"onem": "Orta", "tur": "Yürürlük maddesi yok", "kim": b.id, "aciklama": f"{b.baslik[:80]}: 'yürürlüğe girer' içeren madde bulunamadı"})
    print(f"[OK] {len(basliklar)} başlık · {sum(bool(b.anahtar) for b in basliklar)} anahtar kelime eşleşmesi · {sum(bool(b.metin) for b in basliklar)} tam metin")
    durum = "kullanılmadı (--model-yok)"
    if model:
        hedef = [b for b in basliklar if b.metin]
        llm.onay_al(f"{len(basliklar)} başlık ve şirket profili sınıflandırma için, ilgili çıkan düzenlemelerin tam metinleri (en çok {len(hedef)}) "
                    "özet için gönderilecek. Resmî Gazete metinleri kamuya açıktır; profil dosyanızda gizli bilgi olmadığından emin olun.", evet)
        uy += siniflandir(basliklar, profil)
        for b in basliklar:
            if b.metin and b.ilgi in ("yuksek", "orta"):
                uy += ozetle(b, profil)
        durum = llm.kullanim_ozeti()
    for b in ilgili({"basliklar": basliklar}):
        if not b.metin and b.ilgi in ("yuksek", "orta"):
            uy.append({"onem": "Orta", "tur": "Tam metin gerekli", "kim": b.id, "aciklama": f"{b.baslik[:90]} — ilgili görüldü; metni kaydedip özetletin"})
    bugun = bugun or date.today()
    for b in basliklar:
        for t, _ in b.yururluk:
            if t and 0 <= (t - bugun).days <= 30 and b.ilgi in ("yuksek", "orta", "anahtar"):
                uy.append({"onem": "Yüksek", "tur": "Yakında yürürlük", "kim": b.id, "aciklama": f"{b.baslik[:80]} — {t:%d.%m.%Y} ({(t - bugun).days} gün)"})
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    uy.sort(key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))
    s = {"basliklar": basliklar, "uyarilar": uy, "model": durum, "profil": profil, "bugun": bugun}
    rapor_yaz(cikti, s)
    s["bulten"] = cikti.with_name(cikti.stem + "_bulten.md")
    bulten_yaz(s["bulten"], s)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Resmî Gazete fihristini şirket profiline göre süzer, ilgili düzenlemelerin etkisini özetler.")
    p.add_argument("--fihrist", type=Path, nargs="+", default=[ORNEK / "fihrist_2026-10-09.txt"],
                   help="Resmî Gazete içindekiler sayfası (metin olarak kaydedilmiş .txt); birden çok gün verilebilir")
    p.add_argument("--metinler", type=Path, help="Tam metin klasörü (.txt; ilk satır düzenlemenin başlığı)")
    p.add_argument("--profil", type=Path, default=ORNEK / "profil.txt", help="Şirket profili: Faaliyet, Birimler, Anahtar kelimeler satırları")
    p.add_argument("--ilanlar", action="store_true", help="İlân bölümünü de dahil et")
    p.add_argument("--model-yok", action="store_true", help="Modele gönderme; yalnız fihrist ve anahtar kelime taraması")
    p.add_argument("--bugun", help="Rapor tarihi GG.AA.YYYY (yakın yürürlük uyarısı; örnek veride 09.10.2026)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "mevzuat_ozeti.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    ornek = a.fihrist == [ORNEK / "fihrist_2026-10-09.txt"]
    metinler = a.metinler or (ORNEK / "metinler" if ornek else None)
    for y in (*a.fihrist, a.profil, metinler):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    bugun = None
    if a.bugun:
        try:
            bugun = datetime.strptime(a.bugun, "%d.%m.%Y").date()
        except ValueError:
            print("[X] --bugun GG.AA.YYYY biçiminde olmalı")
            return 2
    elif ornek:
        bugun = date(2026, 10, 9)
    try:
        s = calistir(a.fihrist, a.profil, a.cikti, metin_klasoru=metinler, model=not a.model_yok, ilanlar=a.ilanlar, evet=a.evet, bugun=bugun)
    except (llm.LLMHatasi, ValueError) as h:
        print(f"[X] {h}")
        return 1
    for b in ilgili(s):
        print(f"     [{ILGI_ADI[b.ilgi]}] {b.baslik[:100]}")
    for u in s["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['kim']} · {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[OK] Bülten: {s['bulten'].resolve()}")
    if not a.model_yok:
        print(f"[i] Kullanım: {s['model']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
