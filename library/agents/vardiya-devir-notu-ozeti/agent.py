"""
Vardiya Devir Notu Özeti — Workers / Workless AI Agent
Üretim › Vardiya Amiri

1. Kod, vardiya defterini satır satır numaralar; güvenlik anahtar kelimesi geçen satırları (kaza, yaralanma, ramak kala,
   yangın, duman, kaçak, sızıntı, dökülme, çarpma, düşme, kimyasal, acil durdurma) ve duruş süresi yazılı satırları
   (ör. "Duruş 45 dk") bulur.
2. Model, devir notunu bölümlere ayırır: vardiya bilgisi, üretim durumu, açık sorunlar (öncelik, sonraki vardiyaya
   aksiyon, sorumlu), arıza ve duruşlar, güvenlik, kalite, diğer bilgiler. Her madde defterdeki satır numaralarına
   dayanmak zorundadır; defterde olmayan bilgi, sorumlu ya da süre uydurulmaz.
3. Kod taslağı denetler: her madde geçerli satıra dayanıyor mu, güvenlik satırlarının hepsi güvenlik bölümünde mi,
   duruş satırları arıza / duruş bölümünde mi ve dakikalar defterdekiyle aynı mı, defterde hiçbir bölüme girmeyen
   satır kalmış mı. Sorun varsa taslak modele geri gönderilir (en fazla --tur kez); kalanlar raporda işaretlenir.

Kullanım:
    python agent.py                                              # örnek vardiya defteriyle dener
    python agent.py --girdi defter.docx --tur 3
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import belge
import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
GUVENLIK = ("kaza", "yaralan", "ramak kala", "yangin", "duman", "kivilcim", "kacak", "sizinti", "dokul", "carp", "dustu", "dusme", "kimyasal", "acil durdur", "loto", "elektrik akimina")
SURE = re.compile(r"(\d+(?:[.,]\d+)?)\s*(dk|dakika|saat|sa)\b", re.I)
KAYNAK = {"type": "array", "items": {"type": "integer"}}


def _madde(**alanlar) -> dict:
    alanlar["kaynak_satirlar"] = KAYNAK
    return {"type": "array", "items": {"type": "object", "properties": alanlar, "required": list(alanlar), "additionalProperties": False}}


S = {"type": "string"}
SEMA = {
    "type": "object",
    "properties": {
        "ozet": S,
        "vardiya": {"type": "object", "properties": {"hat_bolum": S, "tarih": S, "vardiya": S, "amir": S, "kaynak_satirlar": KAYNAK},
                    "required": ["hat_bolum", "tarih", "vardiya", "amir", "kaynak_satirlar"], "additionalProperties": False},
        "uretim": _madde(konu=S, detay=S),
        "acik_sorunlar": _madde(konu=S, durum=S, oncelik={"type": "string", "enum": ["Yüksek", "Orta", "Düşük"]}, sonraki_vardiya_aksiyonu=S, sorumlu=S),
        "ariza_ve_duruslar": _madde(ekipman=S, belirti=S, durus_dk={"type": ["number", "null"]}, yapilan=S,
                                    durum={"type": "string", "enum": ["Çözüldü", "Devam ediyor", "Parça bekliyor", "Planlı"]}),
        "guvenlik": _madde(olay=S, yapilan=S, takip=S),
        "kalite": _madde(konu=S, yapilan=S),
        "diger": _madde(konu=S, detay=S),
    },
    "required": ["ozet", "vardiya", "uretim", "acik_sorunlar", "ariza_ve_duruslar", "guvenlik", "kalite", "diger"],
    "additionalProperties": False,
}
BOLUMLER = (("uretim", "Üretim durumu"), ("acik_sorunlar", "Açık sorunlar"), ("ariza_ve_duruslar", "Arıza ve duruşlar"), ("guvenlik", "Güvenlik"),
            ("kalite", "Kalite"), ("diger", "Diğer bilgiler"))


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").translate(_TR).lower()).split())


def satirlar(metin: str) -> list[str]:
    return [s.strip() for s in metin.splitlines() if s.strip()]


def dakika(satir: str) -> float:
    top = 0.0
    for sayi, birim in SURE.findall(satir):
        v = float(sayi.replace(",", "."))
        top += v * 60 if katla(birim) in ("saat", "sa") else v
    return top


def on_inceleme(l: list[str]) -> dict:
    guv = [i for i, s in enumerate(l, 1) if any(k in katla(s) for k in GUVENLIK)]
    durus = {i: dakika(s) for i, s in enumerate(l, 1) if "durus" in katla(s) and dakika(s) > 0}
    return {"guvenlik": guv, "durus": durus}


def denetle(t: dict, l: list[str], on: dict) -> list[str]:
    n = len(l)
    sorunlar, kullanilan = [], set(t["vardiya"]["kaynak_satirlar"])
    for anahtar, ad in BOLUMLER:
        for m in t[anahtar]:
            ks = m["kaynak_satirlar"]
            ilk = next(iter(v for k, v in m.items() if k != "kaynak_satirlar" and isinstance(v, str)), "")
            if not ks:
                sorunlar.append(f"{ad} · '{ilk}': kaynak satır numarası yok")
            gecersiz = [k for k in ks if not 1 <= k <= n]
            if gecersiz:
                sorunlar.append(f"{ad} · '{ilk}': defterde olmayan satır numarası {gecersiz}")
            kullanilan |= set(ks)
    guv_kaynak = {k for m in t["guvenlik"] for k in m["kaynak_satirlar"]}
    for i in on["guvenlik"]:
        if i not in guv_kaynak:
            sorunlar.append(f"Satır {i} güvenlikle ilgili ('{l[i - 1][:70]}') ama Güvenlik bölümünde yok")
    for m in t["ariza_ve_duruslar"]:
        beklenen = sum(on["durus"].get(k, 0) for k in m["kaynak_satirlar"])
        if beklenen and (m["durus_dk"] is None or abs(m["durus_dk"] - beklenen) > 0.5):
            sorunlar.append(f"Arıza ve duruşlar · '{m['ekipman']}': durus_dk {m['durus_dk']}, defterde {beklenen:g} dk yazıyor")
        if not beklenen and m["durus_dk"]:
            sorunlar.append(f"Arıza ve duruşlar · '{m['ekipman']}': defterde süre yazmıyor ama durus_dk {m['durus_dk']} verilmiş (süre uydurulmamalı)")
    ariza_kaynak = {k for m in t["ariza_ve_duruslar"] for k in m["kaynak_satirlar"]}
    for i, dk in on["durus"].items():
        if i not in ariza_kaynak:
            sorunlar.append(f"Satır {i} {dk:g} dk duruş içeriyor ama Arıza ve duruşlar bölümünde yok")
    for i, s in enumerate(l, 1):
        if i not in kullanilan and len(s.split()) >= 3:
            sorunlar.append(f"Satır {i} hiçbir bölümde kullanılmadı: '{s[:70]}'")
    return sorunlar


def uret(l: list[str], on: dict, duzeltme: list[str] | None = None, onceki: dict | None = None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    numarali = "\n".join(f"{i}: {llm.maskele(s)}" for i, s in enumerate(l, 1))
    parcalar = [f"<vardiya_defteri>\n{numarali}\n</vardiya_defteri>",
                "<kod_bulgulari>\nGüvenlik anahtar kelimesi geçen satırlar: " + (", ".join(map(str, on["guvenlik"])) or "yok")
                + "\nDuruş süresi yazılı satırlar: " + (", ".join(f"{i} ({d:g} dk)" for i, d in on["durus"].items()) or "yok") + "\n</kod_bulgulari>"]
    if duzeltme:
        parcalar.append("<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False) + "\nKodun bulduğu sorunlar:\n" + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")
    return llm.json_iste(sistem, "\n".join(parcalar), SEMA)


BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)


def _madde_metni(anahtar: str, m: dict) -> tuple[str, str, str, str]:
    """(başlık, detay, öncelik / durum, sorumlu)"""
    if anahtar == "uretim":
        return m["konu"], m["detay"], "", ""
    if anahtar == "acik_sorunlar":
        return m["konu"], f"{m['durum']} → Sonraki vardiya: {m['sonraki_vardiya_aksiyonu']}", m["oncelik"], m["sorumlu"]
    if anahtar == "ariza_ve_duruslar":
        sure = f" · duruş {m['durus_dk']:g} dk" if m["durus_dk"] else ""
        return m["ekipman"], f"{m['belirti']}{sure} · {m['yapilan']}", m["durum"], ""
    if anahtar == "guvenlik":
        return m["olay"], f"{m['yapilan']} · Takip: {m['takip']}", "", ""
    if anahtar == "kalite":
        return m["konu"], m["yapilan"], "", ""
    return m["konu"], m["detay"], "", ""


def devir_metni(t: dict) -> str:
    v = t["vardiya"]
    out = [f"# Vardiya Devir Notu — {v['hat_bolum']} · {v['tarih']} · {v['vardiya']}", f"Devreden: {v['amir']}", "", t["ozet"]]
    for anahtar, ad in BOLUMLER:
        if t[anahtar]:
            out += ["", f"## {ad}"]
            for m in t[anahtar]:
                b, d, o, s = _madde_metni(anahtar, m)
                out.append(f"- **{b}**{f' [{o}]' if o else ''}: {d}{f' (sorumlu: {s})' if s else ''} _(satır {', '.join(map(str, m['kaynak_satirlar']))})_")
    return "\n".join(out)


def rapor_yaz(cikti: Path, l: list[str], on: dict, t: dict, kalan: list[str], tur: int) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Devir Notu"
    v = t["vardiya"]
    ws.append([f"Vardiya devir notu · {v['hat_bolum']} · {v['tarih']} · {v['vardiya']} · devreden {v['amir']}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([t["ozet"]])
    ws.append([])
    ws.append(["Bölüm", "Madde", "Detay", "Öncelik / Durum", "Sorumlu", "Defter Satırı", "Devralan Onayı"])
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for anahtar, ad in BOLUMLER:
        for m in t[anahtar]:
            b, d, o, s = _madde_metni(anahtar, m)
            ws.append([ad, b, d, o, s, ", ".join(map(str, m["kaynak_satirlar"])), ""])
            if anahtar == "guvenlik" or o in ("Yüksek",):
                ws.cell(ws.max_row, 1).fill = KIRMIZI
            ws.cell(ws.max_row, 7).fill = SARI
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Tüm satırlar bölümlere dağıtıldı, güvenlik ve duruş satırları eksiksiz")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append(["Önemli: Yapay zekâ taslağıdır. Devralan vardiya amiri defterle karşılaştırıp onaylamalıdır; güvenlik olaylarının resmî bildirimi ayrıca yapılır."])
    for j, w in enumerate((18, 28, 70, 14, 14, 10, 12), 1):
        ws.column_dimensions["ABCDEFG"[j - 1]].width = w
    for satir in ws.iter_rows(min_row=4):
        for h in satir:
            h.alignment = UST
    d = wb.create_sheet("Defter")
    d.append(["Satır", "Defter metni", "Kullanıldığı bölüm", "Kod bulgusu"])
    for h in d[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    yer = {}
    for k in t["vardiya"]["kaynak_satirlar"]:
        yer.setdefault(k, []).append("Vardiya bilgisi")
    for anahtar, ad in BOLUMLER:
        for m in t[anahtar]:
            for k in m["kaynak_satirlar"]:
                yer.setdefault(k, []).append(ad)
    for i, s in enumerate(l, 1):
        bulgu = ", ".join(x for x, ok in (("güvenlik", i in on["guvenlik"]), (f"duruş {on['durus'].get(i, 0):g} dk", i in on["durus"])) if ok)
        d.append([i, s, ", ".join(dict.fromkeys(yer.get(i, []))) or "—", bulgu])
        if not yer.get(i) and len(s.split()) >= 3:
            d.cell(d.max_row, 3).fill = KIRMIZI
        if bulgu:
            d.cell(d.max_row, 4).fill = SARI
    for j, w in enumerate((6, 90, 30, 18), 1):
        d.column_dimensions["ABCD"[j - 1]].width = w
    for satir in d.iter_rows(min_row=2):
        for h in satir:
            h.alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    md = cikti.with_suffix(".md")
    md.write_text(devir_metni(t) + "\n", encoding="utf-8")
    return md


def calistir(girdi: Path, cikti: Path, tur: int = 2, evet: bool = False) -> dict:
    l = satirlar(belge.metin_oku(girdi))
    if not l:
        raise belge.BelgeHatasi(f"{girdi.name}: metin bulunamadı.")
    on = on_inceleme(l)
    print(f"[OK] {len(l)} satır · güvenlik satırı {len(on['guvenlik'])} · duruş süresi yazılı satır {len(on['durus'])} ({sum(on['durus'].values()):g} dk)")
    llm.onay_al("Vardiya defteri (numaralı satırlar, telefon / e-posta maskeli) devir notu yazımı için gönderilecek.", evet)
    taslak = uret(l, on)
    kalan = denetle(taslak, l, on)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(l, on, kalan, taslak)
        kalan = denetle(taslak, l, on)
    md = rapor_yaz(cikti, l, on, taslak, kalan, yapilan)
    return {"satirlar": l, "on": on, "taslak": taslak, "kalan": kalan, "tur": yapilan, "md": md}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Vardiya defteri notlarını açık sorunlar, arızalar ve sonraki vardiyaya aksiyonlar olarak özetler.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "vardiya_defteri.txt", help="Vardiya defteri (.txt, .md, .docx, .pdf)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "vardiya_devir_notu.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.cikti, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    t = s["taslak"]
    print(f"[OK] Devir notu hazır · açık sorun {len(t['acik_sorunlar'])} · arıza / duruş {len(t['ariza_ve_duruslar'])} · güvenlik {len(t['guvenlik'])}")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()} (+ {s['md'].name})")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
