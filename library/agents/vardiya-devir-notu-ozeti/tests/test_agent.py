"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri" / "vardiya_defteri.txt"
MESAJLAR = []


def taslak(dogru: bool) -> dict:
    t = {"ozet": "Hat 1 sarıcı rulmanı parça bekliyor; Lot 2610-118 karantinada; silo 3 düşük.",
         "vardiya": {"hat_bolum": "Ekstrüzyon ve Paketleme", "tarih": "13.10.2026", "vardiya": "Gece 00:00-08:00", "amir": "Okan T.", "kaynak_satirlar": [1]},
         "uretim": [{"konu": "Üretim", "detay": "Hat 1 6.850 / 8.000 kg; Hat 2 7.920 / 8.000 kg", "kaynak_satirlar": [11]}],
         "acik_sorunlar": [{"konu": "Hat 1 sarıcı rulmanı", "durum": "Parça yok, talep 26-4471", "oncelik": "Yüksek", "sonraki_vardiya_aksiyonu": "Satın almayı arayın", "sorumlu": "",
                            "kaynak_satirlar": [3, 4]},
                           {"konu": "Silo 3 %18", "durum": "Tanker bekleniyor", "oncelik": "Yüksek", "sonraki_vardiya_aksiyonu": "Lojistikten tanker saatini teyit edin", "sorumlu": "",
                            "kaynak_satirlar": [7]},
                           {"konu": "Personel", "durum": "1 operatör eksik", "oncelik": "Düşük", "sonraki_vardiya_aksiyonu": "Planlamaya bildirin", "sorumlu": "", "kaynak_satirlar": [10]}],
         "ariza_ve_duruslar": [{"ekipman": "Hat 1 sarıcı motoru", "belirti": "Aşırı akım", "durus_dk": 45 if dogru else 40, "yapilan": "Rulman değişimi bekliyor",
                                "durum": "Parça bekliyor", "kaynak_satirlar": [3, 4]},
                               {"ekipman": "Hat 2 ekstrüder", "belirti": "Kafa sıcaklığı alarmı", "durus_dk": None, "yapilan": "Kendiliğinden normale döndü", "durum": "Çözüldü",
                                "kaynak_satirlar": [2]}],
         "guvenlik": [{"olay": "Forklift palet devirdi (ramak kala)", "yapilan": "Bildirim yazıldı", "takip": "İSG değerlendirmesi", "kaynak_satirlar": [6]}],
         "kalite": [{"konu": "Lot 2610-118 alt tolerans", "yapilan": "Karantina", "kaynak_satirlar": [5]}],
         "diger": []}
    if dogru:
        t["ariza_ve_duruslar"].append({"ekipman": "Hat 2 kesici bıçak", "belirti": "Planlı değişim", "durus_dk": 25, "yapilan": "Değiştirildi", "durum": "Planlı", "kaynak_satirlar": [8]})
        t["guvenlik"].append({"olay": "Basınçlı hava kaçağı", "yapilan": "Bakıma bildirildi", "takip": "Sabah bakılacak", "kaynak_satirlar": [9]})
    return t


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return taslak("<duzeltme>" in kullanici)


class OnIncelemeTesti(unittest.TestCase):
    def test_on_inceleme(self):
        l = agent.satirlar(ORNEK.read_text(encoding="utf-8"))
        on = agent.on_inceleme(l)
        self.assertEqual((len(l), on["guvenlik"], on["durus"]), (11, [6, 9], {3: 45.0, 8: 25.0}))
        self.assertEqual([agent.dakika(s) for s in ("Duruş 1,5 saat", "duruş 20 dk + 10 dakika", "yok")], [90.0, 30.0, 0.0])

    def test_denetim(self):
        l = agent.satirlar(ORNEK.read_text(encoding="utf-8"))
        on = agent.on_inceleme(l)
        self.assertEqual(agent.denetle(taslak(True), l, on), [])
        s = agent.denetle(taslak(False), l, on)
        self.assertEqual(len(s), 5)
        self.assertTrue(any("Satır 9 güvenlikle ilgili" in x for x in s) and any("durus_dk 40" in x for x in s) and any("Satır 8 25 dk duruş" in x for x in s))
        t = taslak(True)
        t["kalite"][0]["kaynak_satirlar"] = [99]
        t["ariza_ve_duruslar"][1]["durus_dk"] = 5
        self.assertEqual(len(agent.denetle(t, l, on)), 3)


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "v.xlsx"
            s = agent.calistir(ORNEK, cikti, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            self.assertIn("6: 03:15 Paketleme", MESAJLAR[0])
            self.assertIn("Satır 9 güvenlikle ilgili", MESAJLAR[1].split("<duzeltme>")[1])
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Devir Notu", "Defter"])
            self.assertEqual([wb["Defter"].cell(r, 3).value for r in (2, 10)], ["Vardiya bilgisi", "Güvenlik"])
            self.assertIn("## Güvenlik", s["md"].read_text(encoding="utf-8"))

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(ORNEK, Path(tmp) / "v.xlsx", tur=0, evet=True)
        self.assertEqual(len(s["kalan"]), 5)


if __name__ == "__main__":
    unittest.main()
