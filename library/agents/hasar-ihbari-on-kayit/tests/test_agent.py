"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import os
import re
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

IHBAR = KLASOR / "ornek_veri" / "ihbarlar.csv"
POLICE = KLASOR / "ornek_veri" / "policeler.csv"
MESAJLAR = []


def k(id, police, brans, tarih, il, tur, tutar, sp, kp, tutanak, yar, alintilar=()):
    return {"id": id, "police_no": police, "brans": brans, "hasar_tarihi": tarih, "hasar_saati": "", "il": il, "ilce": "", "yer": "", "hasar_turu": tur,
            "tahmini_tutar": tutar, "sigortali_plaka": sp, "karsi_plaka": kp, "tutanak": tutanak, "yaralanma": yar, "olay_ozeti": "özet",
            "eksik_bilgiler": [], "istenecek_belgeler": [], "alintilar": [{"alan": a, "alinti": b} for a, b in alintilar]}


DOGRU = {
    "IH-2610-001": k("IH-2610-001", "KS-2026-118734", "Kasko", "12.10.2026", "Ankara", "Trafik kazası (çarpışma)", 42500, "34 KLM 4521", "06 ABC 123",
                     "Anlaşmalı kaza tespit tutanağı", False, [("tutar", "yaklaşık 42.500 TL"), ("karsi_plaka", "06 ABC 123 plakalı araç")]),
    "IH-2610-002": k("IH-2610-002", "KN-2025-550912", "Konut", "12.10.2026", "İzmir", "Su / tesisat hasarı", None, "", "", "Belirtilmemiş", False,
                     [("hasar_tarihi", "Dün gece")]),
    "IH-2610-003": k("IH-2610-003", "KS-2026-120045", "Kasko", "06.10.2026", "Kocaeli", "Park hâlinde hasar", 3800, "34 TRZ 778", "", "Yok", False),
    "IH-2610-004": k("IH-2610-004", "IS-2026-077310", "İşyeri", "13.10.2026", "Bursa", "Hırsızlık", 15000, "", "", "Polis / jandarma tutanağı", False),
    "IH-2610-005": k("IH-2610-005", "TR-2026-301122", "Trafik", "10.10.2026", "İstanbul", "Bedeni hasar (yaralanma)", None, "16 BRS 61", "34 DEF 909",
                     "Polis / jandarma tutanağı", True),
    "IH-2610-006": k("IH-2610-006", "", "Kasko", "", "", "Cam kırılması", None, "35 IZM 350", "", "Belirtilmemiş", False),
}


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    d = copy.deepcopy(DOGRU)
    if "<duzeltme>" not in kullanici:
        d["IH-2610-002"]["hasar_tarihi"] = "11.10.2026"
        d["IH-2610-005"]["yaralanma"] = False
        d["IH-2610-006"]["police_no"] = "KS-2026-099001"
    ids = re.findall(r'<ihbar id="([^"]+)"', kullanici.split("<duzeltme>")[0])
    return {"ihbarlar": [d[i] for i in ids]}


class KodTesti(unittest.TestCase):
    def test_yardimcilar(self):
        self.assertEqual(agent.plakalar("arkadan 06 ABC 123 çarptı, benim 34klm4521, tel 0532 111 22 33"), ["06 ABC 123", "34 KLM 4521"])
        self.assertEqual(agent.plaka_norm("16brs61"), "16 BRS 61")
        self.assertEqual(agent.is_gunu_farki(date(2026, 10, 6), date(2026, 10, 14)), 6)
        self.assertEqual(agent.is_gunu_farki(date(2026, 10, 10), date(2026, 10, 14)), 3)
        self.assertTrue(agent.tarih_dayanagi(date(2026, 10, 12), "Dün gece su sızdı", date(2026, 10, 13)))
        self.assertFalse(agent.tarih_dayanagi(date(2026, 10, 11), "Dün gece su sızdı", date(2026, 10, 13)))
        self.assertTrue(agent.tarih_dayanagi(date(2026, 10, 7), "07/10/2026 sabahı", None))

    def test_denetim(self):
        ihb = agent.ihbarlari_oku(IHBAR)
        self.assertEqual(len(ihb), 6)
        self.assertEqual(agent.denetle(ihb, DOGRU), [])
        d = copy.deepcopy(DOGRU)
        d["IH-2610-001"]["tahmini_tutar"] = 45000
        d["IH-2610-001"]["karsi_plaka"] = "06 ABC 124"
        d["IH-2610-001"]["alintilar"][0]["alinti"] = "yaklaşık 45 bin TL"
        d.pop("IH-2610-004")
        s = agent.denetle(ihb, d)
        self.assertEqual(len(s), 4)
        self.assertTrue(any("45000" in x for x in s) and any("06 ABC 124" in x for x in s) and any("IH-2610-004: sonuç" in x for x in s))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "h.xlsx"
            s = agent.calistir(IHBAR, cikti, POLICE, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            self.assertNotIn("0532 111", MESAJLAR[0])
            ilk = MESAJLAR[1].split("<duzeltme>")[1]
            self.assertTrue("çıkarılamıyor" in ilk and "yaralanma false" in ilk and "KS-2026-099001' metinde" in ilk)
            u = s["uyarilar"]
            self.assertEqual({x: len(v) for x, v in u.items()}, {"IH-2610-001": 0, "IH-2610-002": 1, "IH-2610-003": 1, "IH-2610-004": 0,
                                                                  "IH-2610-005": 1, "IH-2610-006": 2})
            self.assertIn("poliçe süresi dışında", u["IH-2610-002"][0])
            self.assertIn("6 iş günü", u["IH-2610-003"][0])
            self.assertEqual(s["oneriler"]["IH-2610-006"], "KS-2026-099001")
            ws = load_workbook(cikti)["Ön Kayıt"]
            self.assertEqual((ws.cell(2, 1).value, ws.cell(2, 4).value), ("IH-2610-005", "Öncelikli"))

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(IHBAR, Path(tmp) / "h.xlsx", tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 3)
            self.assertEqual(s["oneriler"]["IH-2610-001"], "")
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--girdi", str(IHBAR.with_name("yok.csv")), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
