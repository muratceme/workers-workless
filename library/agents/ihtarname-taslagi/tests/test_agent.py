"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

OLAY = KLASOR / "ornek_veri" / "olay.txt"
TARAF = KLASOR / "ornek_veri" / "taraflar.txt"
TALEP = "Ödenmeyen kira bedellerinin ödenmesi; ödenmezse kira sözleşmesinin feshi ve tahliye"
MESAJLAR = []


def taslak(dogru=True) -> dict:
    a = [{"metin": "Müvekkil {KESIDECI}, konutunu 01.03.2025 tarihli kira sözleşmesiyle muhatap {MUHATAP} adlı kişiye kiraya vermiştir. "
                   "Aylık kira bedeli 25.000 TL olup her ayın 5'ine kadar ödenecektir.", "kaynak_satirlar": [1, 2]},
         {"metin": "Muhatap; Temmuz 2026, Ağustos 2026 ve Eylül 2026 aylarına ait toplam 75.000 TL kira bedelini ödememiş, "
                   "12.09.2026 tarihli hatırlatmaya rağmen ödeme yapmamıştır.", "kaynak_satirlar": [3, 4, 5]}]
    si = ("Yukarıda açıklanan nedenlerle, Temmuz 2026, Ağustos 2026 ve Eylül 2026 aylarına ait toplam 75.000 TL kira bedelini işbu "
          "ihtarnamenin tebliğinden itibaren 30 (otuz) gün içinde ödemenizi; aksi hâlde TBK 315 uyarınca kira sözleşmesinin feshedileceğini "
          "ve tahliye için yasal yollara başvurulacağını ihtaren bildiririz.")
    if not dogru:
        a[1]["metin"] = a[1]["metin"].replace("75.000 TL", "80.000 TL") + " Bu durum TBK 299 uyarınca açık bir ihlaldir; muhatap dolandırıcı gibi davranmaktadır."
        a[0]["metin"] = a[0]["metin"].replace("muhatap {MUHATAP} adlı kişiye", "{MUHATAP}'a")
        si = "Borcun 15 gün içinde ödenmesini ihtar ederiz."
    return {"konu": "Ödenmeyen kira bedellerinin ödenmesi ihtarı", "aciklamalar": a, "sonuc_ve_istem": si,
            "eksik_bilgiler": ["Keşidecinin banka hesap (IBAN) bilgisi"], "ic_not": "Sözleşmedeki ödeme günü teyit edilmeli."}


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return taslak("<duzeltme>" in kullanici)


class KodTesti(unittest.TestCase):
    def test_gizleme(self):
        t = agent.taraflari_oku(TARAF)
        l = agent.satirlar(agent.gizle(OLAY.read_text(encoding="utf-8"), t))
        self.assertEqual(len(l), 6)
        self.assertIn("{KESIDECI}", l[0])
        self.assertIn("{MUHATAP}'e", l[0])

    def test_denetim(self):
        t = agent.taraflari_oku(TARAF)
        l = agent.satirlar(agent.gizle(OLAY.read_text(encoding="utf-8"), t))
        self.assertEqual(agent.denetle(taslak(), l, TALEP, 30, "kira-konut", ""), [])
        s = agent.denetle(taslak(False), l, TALEP, 30, "kira-konut", "")
        self.assertEqual(len(s), 7, s)
        self.assertTrue(any("80000" in x for x in s) and any("madde 299" in x for x in s) and any("'30 gün'" in x for x in s)
                        and any("feshed" in x for x in s) and any("ek getirilmiş" in x for x in s) and any("dolandırıcı" in x for x in s))
        # TBK 315 genel ihtarda izinli değil; dayanakla verilirse izinli
        self.assertTrue(any("madde 315" in x for x in agent.denetle(taslak(), l, TALEP, 30, "genel", "")))
        self.assertEqual(agent.denetle(taslak(), l, TALEP, 30, "genel", "TBK 315"), [])


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "i.xlsx"
            s = agent.calistir(OLAY, TARAF, TALEP, 30, cikti, "kira-konut", evet=True, tarih=date(2026, 10, 15))
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            for gizli in ("Ayşe", "Yılmaz", "Can Demir", "Deneme Sok"):
                self.assertNotIn(gizli, MESAJLAR[0])
            m = cikti.with_suffix(".txt").read_text(encoding="utf-8")
            self.assertIn("MUHATAP       : Can Demir", m)
            self.assertIn("muhatap Can Demir adlı kişiye", m)
            self.assertIn("tebliğ şerhli bir nüshasının", m)
            self.assertEqual(load_workbook(cikti).sheetnames, ["Paragraflar", "Eksik Bilgiler"])

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_sure_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaises(llm.LLMHatasi):
                agent.calistir(OLAY, TARAF, TALEP, 15, Path(tmp) / "i.xlsx", "kira-konut", evet=True)
            s = agent.calistir(OLAY, TARAF, TALEP, 30, Path(tmp) / "i.xlsx", "kira-konut", tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 7)
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--evet", "--sure", "10"]), 1)


if __name__ == "__main__":
    unittest.main()
