"""
Kaza ve Ramak Kala Bildirimlerinden Olay Kaydı — Gerisi Bende AI Agent
İş Sağlığı ve Güvenliği › İşyeri Hekimi

Çalışan ve amirlerin serbest metin kaza / ramak kala bildirimlerini "İş Kazası ve Ramak Kala Analizi" kod bloğunun olay
kaydına dönüştürür; iş kazalarında yaralanma bilgisini ve SGK bildirim son gününü ekler.
1. Kod bildirimleri okur ("Bildirim: No | Tarih: GG.AA.YYYY | Saat: SS:DD | Bölüm: … | Bildiren: Ad Soyad" başlıklı
   bloklar). Bildiren ve --gizle adları takma adla, telefon, e-posta, TCKN maskelenir.
2. Model türü (İş Kazası / Ramak Kala / Belirsiz) kanıt cümlesiyle, neden kategorisini, kayıp günü, vardiyayı, kısa
   açıklamayı ve yaralanma bilgisini çıkarır.
3. Kod denetler: kanıt bildirimde aynen var mı; bildirimde zarar ifadesi (yaralanma, sıkışma, revir, sevk, rapor …)
   varken "Ramak Kala", zarar yazmazken "İş Kazası" denmiş mi; kayıp gün ve sayılar bildirimde geçiyor mu; neden
   listede mi; açıklamada kişi adı var mı. Sorun varsa yalnız o bildirim modele geri gönderilir.
4. Kod iş kazaları için SGK bildirim son gününü hesaplar (kaza günü sayılmaz, hafta sonu ve --tatiller hariç 3 iş günü).

Kullanım:
    python agent.py                                                    # örnek 4 bildirim
    python agent.py --bildirimler bildirimler.docx --tatiller 29.10.2026 --gizle "Ad Soyad"
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}
ZARAR = re.compile(r"\b(yaralan|kesil|burk|kirik|kirild|incin|tutuldu|tutuldugu|sikisti|sikisma|ezil|yanik|yandi|carpti|agri|kanama|dikis|revir|ilk yardim|"
                   r"hastane|sevk|rapor verildi|is goremez|bayil|zehirlen|cizil|morar)\w*")
ZARARSIZ = re.compile(r"\b(kimse yaralanmadi|yaralanan olmadi|temas olmadi|temas etmedi|zarar gormedi|yaralanma yok)\b")
NEDENLER = ["El aleti kullanımı", "Kayma ve düşme", "Forklift ve araç hareketi", "Elektrik", "Makine ve ekipman", "Makine koruyucusu eksikliği",
            "Malzeme düşmesi", "Elle taşıma", "Kimyasal", "Diğer"]


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def sayilar(metin: str) -> set[str]:
    temiz = re.sub(r"\b\d{1,2}\.\d{1,2}\.\d{4}\b|\b\d{1,2}:\d{2}\b|\b[A-Z]{1,3}-[\w-]+\b", " ", str(metin))
    return {str(int(x)) if x.isdigit() else x.replace(",", ".") for x in re.findall(r"\d+(?:[.,]\d+)?", temiz)}


def tarih_coz(x) -> date | None:
    m = re.search(r"\b(\d{1,2})\.(\d{1,2})\.(\d{4})\b", str(x or ""))
    try:
        return date(int(m[3]), int(m[2]), int(m[1])) if m else None
    except ValueError:
        return None


def is_gunu_ekle(gun: date, n: int, tatiller: set[date]) -> date:
    """Olay günü sayılmaz; cumartesi, pazar ve verilen resmî tatiller iş günü değildir."""
    while n:
        gun += timedelta(days=1)
        if gun.weekday() < 5 and gun not in tatiller:
            n -= 1
    return gun


# ----------------------------------------------------------------------------
# Girdi
# ----------------------------------------------------------------------------

def takma_adlar(adlar: list[str]) -> dict[str, str]:
    return {a: f"[KİŞİ-{i}]" for i, a in enumerate(dict.fromkeys(x.strip() for x in adlar if x.strip()), 1)}


def maskele(metin: str, kisiler: dict[str, str]) -> str:
    for g, t in sorted(kisiler.items(), key=lambda x: -len(x[0])):
        metin = re.sub(re.escape(g), t, metin, flags=re.I)
    return llm.maskele(metin)


def bildirimleri_oku(yol: Path, gizle: list[str] | None = None) -> list[dict]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: bildirim dosyası bulunamadı.")
    out, son = [], None
    for s in belge.metin_oku(yol).replace("\r", "").splitlines():
        m = re.match(r"^\s*Bildirim\s*:\s*(.+)$", s, re.I)
        if m:
            alan = [x.strip() for x in m[1].split("|")]
            ek = {katla(a.split(":", 1)[0]): a.split(":", 1)[1].strip() for a in alan[1:] if ":" in a}
            saat = re.fullmatch(r"(\d{1,2})[:.](\d{2})", ek.get("saat", ""))
            son = {"no": alan[0], "tarih": tarih_coz(ek.get("tarih")), "saat": f"{int(saat[1]):02d}:{saat[2]}" if saat else "", "bolum": ek.get("bolum", ""),
                   "bildiren": ek.get("bildiren", ""), "metin": ""}
            out.append(son)
        elif son and s.strip():
            son["metin"] = (son["metin"] + " " + s.strip()).strip()
    out = [b for b in out if b["metin"]]
    if not out:
        raise llm.LLMHatasi(f"{yol.name}: 'Bildirim: No | Tarih: GG.AA.YYYY | Saat: SS:DD | Bölüm: …' başlıklı bildirim yok.")
    tekrar = sorted({b["no"] for b in out if [x["no"] for x in out].count(b["no"]) > 1})
    if tekrar:
        raise llm.LLMHatasi(f"Aynı bildirim numarası birden çok kez geçiyor: {', '.join(tekrar)}")
    kisiler = takma_adlar([b["bildiren"] for b in out] + list(gizle or []))
    for b in out:
        b["metin_m"] = maskele(b["metin"], kisiler)
    return out


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

def sema(nedenler: list[str]) -> dict:
    return {"type": "object", "properties": {
        "tur": {"type": "string", "enum": ["İş Kazası", "Ramak Kala", "Belirsiz"]}, "kanit": S, "neden": {"type": "string", "enum": nedenler},
        "kayip_gun": S, "vardiya": S, "aciklama": S,
        "yaralanma": {"type": "object", "properties": {"vucut_bolgesi": S, "yaralanma_turu": S, "mudahale": S,
                                                       "sevk": {"type": "string", "enum": ["Hastaneye sevk", "İşyerinde tedavi", "Yok", "Belirsiz"]}},
                      "required": ["vucut_bolgesi", "yaralanma_turu", "mudahale", "sevk"], "additionalProperties": False},
        "eksikler": {"type": "array", "items": S}},
        "required": ["tur", "kanit", "neden", "kayip_gun", "vardiya", "aciklama", "yaralanma", "eksikler"], "additionalProperties": False}


def denetle(b: dict, c: dict, nedenler: list[str]) -> list[str]:
    no, km, sorun = b["no"], katla(b["metin_m"]), []
    kk = katla(c["kanit"])
    if len(kk.split()) < 2 or kk not in km:
        sorun.append(f"{no}: kanıt bildirimde aynen geçmiyor; türü belirleyen cümle parçasını kelimesi kelimesine yazın")
    zarar = [m[0] for m in ZARAR.finditer(km)]
    zararsiz = bool(ZARARSIZ.search(km))
    if c["tur"] == "Ramak Kala" and zarar and not zararsiz:
        sorun.append(f"{no}: bildirimde zarar ifadesi var ({', '.join(dict.fromkeys(zarar))}); çalışan zarar gördüyse hafif de olsa 'İş Kazası'dır")
    if c["tur"] == "İş Kazası" and (not zarar or zararsiz):
        sorun.append(f"{no}: bildirimde zarar / yaralanma ifadesi yok; 'Ramak Kala' ya da 'Belirsiz' olmalı")
    if c["neden"] not in nedenler:
        sorun.append(f"{no}: neden '{c['neden']}' listede yok")
    if c["kayip_gun"].strip():
        if not re.fullmatch(r"\d+", c["kayip_gun"].strip()):
            sorun.append(f"{no}: kayıp gün yalnız sayı olmalı ('{c['kayip_gun']}')")
        elif c["kayip_gun"].strip() != "0" and str(int(c["kayip_gun"])) not in sayilar(b["metin"]):
            sorun.append(f"{no}: kayıp gün {c['kayip_gun']} bildirimde yazmıyor")
        if c["tur"] == "Ramak Kala" and c["kayip_gun"].strip() not in ("", "0"):
            sorun.append(f"{no}: ramak kalada kayıp gün olamaz")
    if c["vardiya"] and katla(c["vardiya"]).split()[0] not in km:
        sorun.append(f"{no}: vardiya '{c['vardiya']}' bildirimde yazmıyor")
    metin = " ".join([c["aciklama"], *c["yaralanma"].values()])
    fazla = sorted(sayilar(metin) - sayilar(b["metin"]), key=float)
    if fazla:
        sorun.append(f"{no}: bildirimde olmayan sayı: {', '.join(fazla)}")
    if re.search(r"\[KİŞİ-\d+\]", c["aciklama"]):
        sorun.append(f"{no}: açıklamada kişi adı var; olay kaydı kişi adı içermemeli")
    if len(c["aciklama"]) > 120:
        sorun.append(f"{no}: açıklama {len(c['aciklama'])} karakter; en fazla 120")
    if c["tur"] == "Ramak Kala" and any(c["yaralanma"][k] for k in ("vucut_bolgesi", "yaralanma_turu", "mudahale")):
        sorun.append(f"{no}: ramak kalada yaralanma bilgisi olamaz")
    if c["tur"] == "İş Kazası" and not (c["yaralanma"]["vucut_bolgesi"] or c["yaralanma"]["yaralanma_turu"]):
        sorun.append(f"{no}: iş kazasında bildirimde yazan vücut bölgesi / yaralanma türü boş")
    return sorun


def uret(bildirimler: list[dict], nedenler: list[str], duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    sonuc = dict(onceki or {})
    hedef = bildirimler if duzeltme is None else [b for b in bildirimler if any(s.startswith(b["no"] + ":") for s in duzeltme)]
    for b in hedef:
        mesaj = (f"<nedenler>{' | '.join(nedenler)}</nedenler>\n<bildirim no=\"{b['no']}\" tarih=\"{b['tarih']:%d.%m.%Y}\" saat=\"{b['saat']}\" bolum=\"{b['bolum']}\">"
                 if b["tarih"] else f"<nedenler>{' | '.join(nedenler)}</nedenler>\n<bildirim no=\"{b['no']}\" saat=\"{b['saat']}\" bolum=\"{b['bolum']}\">")
        mesaj += f"{b['metin_m']}</bildirim>"
        d = [s.split(": ", 1)[1] for s in (duzeltme or []) if s.startswith(b["no"] + ":")]
        if d:
            mesaj += ("\n<duzeltme>\nÖnceki cevap: " + json.dumps(sonuc.get(b["no"]), ensure_ascii=False) + "\nKodun bulduğu sorunlar:\n"
                      + "\n".join(f"- {s}" for s in d) + "\n</duzeltme>")
        sonuc[b["no"]] = llm.json_iste(sistem, mesaj, sema(nedenler))
    return sonuc


def duzeltme_dongusu(uret_f, denetle_f, tur: int) -> tuple[dict, list[str], int]:
    taslak = uret_f(None, None)
    kalan = denetle_f(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret_f(kalan, taslak)
        kalan = denetle_f(taslak)
    return taslak, kalan, yapilan


# ----------------------------------------------------------------------------
# Çıktı
# ----------------------------------------------------------------------------

def tablo(ws, bas: list[str], satirlar: list[list], ai: set[int] = frozenset(), genis: dict | None = None) -> None:
    ws.append(bas)
    for h in ws[1]:
        h.fill, h.font, h.alignment = BASLIK_DOLGU, BASLIK_YAZI, UST
    for r in satirlar:
        ws.append(r)
        for j, v in enumerate(r, 1):
            c = ws.cell(ws.max_row, j)
            c.alignment = UST
            if isinstance(v, date):
                c.number_format = "DD.MM.YYYY"
            if j in ai and v not in (None, ""):
                c.fill = AI_DOLGU
    for j in range(1, len(bas) + 1):
        ws.column_dimensions[get_column_letter(j)].width = (genis or {}).get(j, 14)
    ws.freeze_panes = "A2"


def sgk_son_gun(b: dict, c: dict, tatiller: set[date]) -> date | None:
    return is_gunu_ekle(b["tarih"], 3, tatiller) if c["tur"] == "İş Kazası" and b["tarih"] else None


def olay_satirlari(bildirimler: list[dict], sonuc: dict) -> list[list]:
    return [[b["tarih"], b["saat"], sonuc[b["no"]]["tur"], b["bolum"], sonuc[b["no"]]["neden"],
             int(sonuc[b["no"]]["kayip_gun"]) if re.fullmatch(r"\d+", sonuc[b["no"]]["kayip_gun"].strip()) else None, sonuc[b["no"]]["vardiya"],
             sonuc[b["no"]]["aciklama"]] for b in bildirimler if sonuc[b["no"]]["tur"] != "Belirsiz"]


def rapor_yaz(cikti: Path, bildirimler: list[dict], sonuc: dict, kalan: list[str], tur: int, tatiller: set[date], bugun: date) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Bildirimler"
    rr = []
    for b in bildirimler:
        c = sonuc[b["no"]]
        y = c["yaralanma"]
        sg = sgk_son_gun(b, c, tatiller)
        durum = "" if not sg else ("Süresi geçti" if bugun > sg else f"{(sg - bugun).days} gün kaldı")
        sor = "; ".join(x.split(": ", 1)[1] for x in kalan if x.startswith(b["no"] + ":"))
        rr.append([b["no"], b["tarih"], b["saat"], b["bolum"], c["tur"], c["neden"], c["kayip_gun"], c["vardiya"], c["aciklama"], y["vucut_bolgesi"],
                   y["yaralanma_turu"], y["mudahale"], y["sevk"], sg, durum, c["kanit"], "\n".join(c["eksikler"]), sor or "Geçti"])
    tablo(ws, ["Bildirim No", "Tarih", "Saat", "Bölüm", "Tür", "Neden", "Kayıp Gün", "Vardiya", "Açıklama", "Vücut Bölgesi", "Yaralanma Türü",
               "Müdahale", "Sevk", "SGK Bildirim Son Günü", "SGK Durumu", "Kanıt", "Eksik Bilgiler", "Son Kontrol"], rr,
          ai={5, 6, 7, 8, 9, 10, 11, 12, 13, 16, 17}, genis={9: 40, 12: 30, 16: 40, 17: 40, 18: 30})
    for n, r in enumerate(rr, 2):
        ws.cell(n, 5).fill = {"İş Kazası": KIRMIZI, "Ramak Kala": SARI}.get(r[4], AI_DOLGU)
        if r[14]:
            ws.cell(n, 15).fill = KIRMIZI if r[14] == "Süresi geçti" else SARI
        ws.cell(n, 18).fill = YESIL if r[17] == "Geçti" else KIRMIZI
    k = wb.create_sheet("Kontrol")
    k.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    k.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    k.cell(2, 1).fill = KIRMIZI if kalan else YESIL
    k.append(["SGK bildirim son günü: kaza günü sayılmadan 3 iş günü (cumartesi, pazar ve --tatiller ile verilen resmî tatiller hariç); başlangıç "
              "olarak kaza tarihi alınır. Kazanın işveren kontrolü dışında olduğu ve sonradan öğrenildiği durumlar ile güncel mevzuat için SGK "
              "düzenlemesini kontrol edin. Tür ve yaralanma bilgileri modelindir (mor); kayda geçmeden önce okuyun. 'Belirsiz' bildirimler olay "
              "dosyasına yazılmaz. Sağlık bilgisi özel nitelikli kişisel veridir; dosyayı yetkisiz kişilerle paylaşmayın."])
    k.column_dimensions["A"].width = 140
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    o = Workbook()
    tablo(o.active, ["Tarih", "Saat", "Tür", "Bölüm", "Neden", "Kayıp Gün", "Vardiya", "Açıklama"], olay_satirlari(bildirimler, sonuc), genis={5: 26, 8: 60})
    y = cikti.with_name(f"{cikti.stem}_olaylar.xlsx")
    o.save(y)
    return y


def calistir(yol: Path, cikti: Path, nedenler: list[str], tatiller: set[date] | None = None, gizle: list[str] | None = None, tur: int = 2,
             evet: bool = False, bugun: date | None = None) -> dict:
    bildirimler = bildirimleri_oku(yol, gizle)
    print(f"[OK] {len(bildirimler)} bildirim okundu")
    llm.onay_al(f"{len(bildirimler)} bildirim metni gönderilecek (bildiren ve --gizle adları takma adla, telefon, e-posta, TCKN maskeli). "
                "Metinde sağlık bilgisi olabilir; gizliyse yerel model (Ollama) kullanın.", evet)
    sonuc, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(bildirimler, nedenler, d, o),
                                             lambda x: [s for b in bildirimler for s in denetle(b, x[b["no"]], nedenler)], tur)
    olay = rapor_yaz(cikti, bildirimler, sonuc, kalan, yapilan, tatiller or set(), bugun or date.today())
    return {"bildirimler": bildirimler, "sonuc": sonuc, "kalan": kalan, "tur": yapilan, "olaylar": olay}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Serbest metin kaza / ramak kala bildirimlerini olay kaydına dönüştürür; SGK bildirim son gününü hesaplar.")
    p.add_argument("--bildirimler", type=Path, default=ORNEK / "bildirimler.txt", help="Bildirimler (.txt/.docx/.pdf); her blok 'Bildirim: No | Tarih: … | Saat: … | Bölüm: … | Bildiren: …' satırıyla başlar")
    p.add_argument("--nedenler", type=Path, default=ORNEK / "neden_kategorileri.txt", help="Neden kategorileri, satır başına bir tane")
    p.add_argument("--tatiller", nargs="*", default=[], metavar="GG.AA.YYYY", help="SGK süresinde sayılmayacak resmî tatil günleri")
    p.add_argument("--gizle", nargs="*", default=[], metavar="AD", help="Metinde geçen ve gönderilmeyecek kazazede / tanık adları")
    p.add_argument("--tarih", help="Bugünün tarihi GG.AA.YYYY (SGK süresinin kalanı için; varsayılan bugün)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kaza_bildirimleri.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    ornek = a.bildirimler == ORNEK / "bildirimler.txt"
    gizle = a.gizle or (["Mehmet Kaya", "Ali Demir", "Zeynep Aksoy"] if ornek else [])
    tatiller = {t for x in a.tatiller if (t := tarih_coz(x))}
    try:
        if not a.nedenler.exists():
            raise llm.LLMHatasi(f"{a.nedenler}: neden kategorileri dosyası bulunamadı.")
        nedenler = [x.strip() for x in a.nedenler.read_text(encoding="utf-8-sig").splitlines() if x.strip()] or NEDENLER
        if "Diğer" not in nedenler:
            nedenler.append("Diğer")
        bugun = tarih_coz(a.tarih) if a.tarih else (date(2026, 10, 16) if ornek else None)
        s = calistir(a.bildirimler, a.cikti, nedenler, tatiller, gizle, a.tur, a.evet, bugun)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    for b in s["bildirimler"]:
        c = s["sonuc"][b["no"]]
        sg = sgk_son_gun(b, c, tatiller)
        print(f"[OK] {b['no']} · {c['tur']} · {c['neden']}" + (f" · SGK bildirim son günü {sg:%d.%m.%Y}" if sg else ""))
    for x in s["kalan"]:
        print(f"[!] Kalan: {x}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[OK] Kod bloğu girdisi: {s['olaylar'].name}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
