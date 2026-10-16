"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import json
import os
import re
import sys
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri" / "bildirimler.txt"
CEVAP = json.loads((Path(__file__).parent / "sahte_cevaplar.json").read_text(encoding="utf-8"))
ADLAR = ["Mehmet Kaya", "Ali Demir", "Zeynep Aksoy"]
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    """B-103'te ilk turda bel tutulmasını 'Ramak Kala' sayar."""
    MESAJLAR.append(kullanici)
    no = re.search(r'<bildirim no="([^"]+)"', kullanici)[1]
    return copy.deepcopy(CEVAP[f"{no}_hatali" if no == "B-103" and "<duzeltme>" not in kullanici else no])


def bildirimler():
    return {b["no"]: b for b in agent.bildirimleri_oku(ORNEK, ADLAR)}


class KodTesti(unittest.TestCase):
    def test_okuma(self):
        b = bildirimler()
        self.assertEqual(list(b), ["B-101", "B-102", "B-103", "B-104"])
        self.assertEqual((b["B-102"]["tarih"], b["B-102"]["saat"], b["B-102"]["bolum"]), (date(2026, 10, 15), "23:40", "Üretim"))
        for x in b.values():
            for ad in ADLAR + ["Kemal Uçar", "Serkan Polat", "Derya Kurt"]:
                self.assertNotIn(ad, x["metin_m"])

    def test_is_gunu(self):
        self.assertEqual(agent.is_gunu_ekle(date(2026, 10, 15), 3, set()), date(2026, 10, 20))     # perşembe → salı
        self.assertEqual(agent.is_gunu_ekle(date(2026, 10, 16), 3, set()), date(2026, 10, 21))     # cuma → çarşamba
        self.assertEqual(agent.is_gunu_ekle(date(2026, 10, 17), 3, set()), date(2026, 10, 21))     # cumartesi → çarşamba
        self.assertEqual(agent.is_gunu_ekle(date(2026, 10, 27), 3, {date(2026, 10, 29)}), date(2026, 11, 2))

    def test_denetim(self):
        b = bildirimler()
        n = agent.NEDENLER
        self.assertEqual([agent.denetle(b[k], CEVAP[k], n) for k in b], [[], [], [], []])
        self.assertEqual(agent.denetle(b["B-103"], CEVAP["B-103_hatali"], n),
                         ["B-103: bildirimde zarar ifadesi var (tutuldugunu, revirde); çalışan zarar gördüyse hafif de olsa 'İş Kazası'dır"])
        c = copy.deepcopy(CEVAP["B-104"]) | {"tur": "İş Kazası", "kayip_gun": "3", "vardiya": "Gece", "kanit": "şişe kırıldı"}
        s = agent.denetle(b["B-104"], c, n)
        for p in ("zarar / yaralanma ifadesi yok", "kayıp gün 3 bildirimde yazmıyor", "vardiya 'Gece'", "kanıt bildirimde aynen geçmiyor", "vücut bölgesi"):
            self.assertTrue(any(p in x for x in s), p)
        c = copy.deepcopy(CEVAP["B-102"])
        c["aciklama"] = "[KİŞİ-5] parmağını 3 yerinden sıkıştırdı" + " ." * 60
        c["neden"] = "Dikkatsizlik"
        s = agent.denetle(b["B-102"], c, n)
        for p in ("olmayan sayı: 3", "kişi adı", "en fazla 120", "listede yok"):
            self.assertTrue(any(p in x for x in s), p)


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ornek_ve_duzeltme(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "k.xlsx"
            self.assertEqual(agent.main(["--evet", "--cikti", str(cikti)]), 0)
            self.assertEqual(len(MESAJLAR), 5)                          # 4 bildirim + B-103 düzeltmesi
            for m in MESAJLAR:
                for ad in ADLAR:
                    self.assertNotIn(ad, m)
            r = [[c.value for c in x] for x in load_workbook(cikti)["Bildirimler"].iter_rows(min_row=2)]
            self.assertEqual([(x[4], x[13], x[14]) for x in r], [("Ramak Kala", None, None), ("İş Kazası", datetime(2026, 10, 20), "4 gün kaldı"),
                                                                  ("İş Kazası", datetime(2026, 10, 21), "5 gün kaldı"), ("Ramak Kala", None, None)])
            o = [[c.value for c in x] for x in load_workbook(Path(tmp) / "k_olaylar.xlsx").active.iter_rows()]
            self.assertEqual(o[0], ["Tarih", "Saat", "Tür", "Bölüm", "Neden", "Kayıp Gün", "Vardiya", "Açıklama"])
            self.assertEqual(o[2][:7], [datetime(2026, 10, 15), "23:40", "İş Kazası", "Üretim", "Makine ve ekipman", 5, "Gece"])

    def test_belirsiz_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(agent.main(["--bildirimler", str(Path(tmp) / "yok.txt"), "--evet"]), 1)
        s = {"B-1": CEVAP["B-101"] | {"tur": "Belirsiz"}}
        self.assertEqual(agent.olay_satirlari([{"no": "B-1", "tarih": None, "saat": "", "bolum": ""}], s), [])   # belirsiz olay dosyasına yazılmaz


if __name__ == "__main__":
    unittest.main()
