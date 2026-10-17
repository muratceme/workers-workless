"""
Ürün Değerlendirme Analizi — Gerisi Bende AI Agent
Müşteri Deneyimi › Müşteri Hizmetleri Temsilcisi

1. Kod yorumları okur; ürün güvenliği belirtilerini (duman, yanık, alev, patlama, elektrik çarpması, kızarıklık, kaşıntı,
   alerji, zehirlenme, kesik) ve kargo / paketleme ifadelerini işaretler.
2. Model her yorumu genel duyguya ve konulara (ürün kalitesi, beden / kalıp, kargo / teslimat, paketleme, fiyat /
   performans, açıklamaya uygunluk, satıcı / müşteri hizmetleri, ürün güvenliği) ayırır; her konu için duygu ve
   yorumdan aynen alıntı verir, tek cümlelik Türkçe özet yazar.
3. Kod denetler: alıntılar yorumda var mı; güvenlik belirtisi olan yorumda "Ürün güvenliği" konusu, kargo / paketleme
   ifadesi olan yorumda ilgili konu var mı. Sorun varsa modele geri gönderilir (en fazla --tur kez).
4. Kod özetler: ürün bazında yorum sayısı, ortalama puan, konu bazında olumlu / olumsuz / net skor, en çok şikâyet
   edilen konu ve örnek alıntılar. Ürün güvenliği yorumları "Acil" sayfasına alınır; puanla duygusu çelişen yorumlar
   işaretlenir.

Kullanım:
    python agent.py                                    # örnek yorumlarla dener
    python agent.py --girdi yorumlar.xlsx
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
PAKET = 30
GUVENLIK_K, KARGO_K, PAKET_K = "Ürün güvenliği", "Kargo / teslimat", "Paketleme"
KONULAR = ["Ürün kalitesi", "Beden / kalıp", KARGO_K, PAKET_K, "Fiyat / performans", "Açıklamaya uygunluk", "Satıcı / müşteri hizmetleri", GUVENLIK_K, "Diğer"]
DUYGU = ["Olumlu", "Nötr", "Olumsuz"]
GUVENLIK = re.compile(r"\b(duman|yandi|yanik|alev|patladi|kivilcim|elektrik carp|carpti elektrik|kizarik|kasinti|alerji|zehirlen|kesti|kesik|bogul)")
KARGO = re.compile(r"\b(kargo|teslimat|kurye|gunde geldi|gec geldi|ertesi gun geldi)")
PAKETLEME = re.compile(r"\b(kutu|kutusu|paket|ambalaj)")
SEMA = {
    "type": "object",
    "properties": {"yorumlar": {"type": "array", "items": {"type": "object", "properties": {
        "id": S, "duygu": {"type": "string", "enum": DUYGU}, "ozet": S,
        "konular": {"type": "array", "items": {"type": "object", "properties": {"konu": {"type": "string", "enum": KONULAR}, "duygu": {"type": "string", "enum": DUYGU}, "alinti": S},
                                               "required": ["konu", "duygu", "alinti"], "additionalProperties": False}}},
        "required": ["id", "duygu", "ozet", "konular"], "additionalProperties": False}}},
    "required": ["yorumlar"],
    "additionalProperties": False,
}


def yorumlari_oku(yol: Path) -> list[dict]:
    r = kayitlar(yol, {"id": ("Yorum No", "No", "ID"), "tarih": ("Tarih", "Yorum Tarihi"), "urun": ("Ürün Kodu", "SKU", "Stok Kodu"), "ad": ("Ürün Adı", "Ürün"),
                       "puan": ("Puan", "Yıldız", "Rating"), "yorum": ("Yorum", "Değerlendirme", "Metin")}, ("yorum",))
    out = []
    for i, x in enumerate(r, 1):
        if not str(x["yorum"]).strip():
            continue
        try:
            puan = float(sayi_norm(x["puan"])) if str(x["puan"]).strip() else None
        except ValueError:
            puan = None
        out.append({"id": str(x["id"] or f"Y{i}"), "tarih": tarih_coz(x["tarih"]), "urun": str(x["urun"] or x["ad"]), "ad": str(x["ad"]), "puan": puan,
                    "yorum": str(x["yorum"]).strip()})
    return out


def isaretler(metin: str) -> list[str]:
    k = katla(metin)
    return [ad for ad, d in ((GUVENLIK_K, GUVENLIK), (KARGO_K, KARGO), (PAKET_K, PAKETLEME)) if d.search(k)]


def puan_yaz(y: dict) -> str:
    return "" if y["puan"] is None else sayi_yaz(y["puan"])


def denetle(yorum: list[dict], sonuc: dict) -> list[str]:
    sorun = []
    for y in yorum:
        s = sonuc.get(y["id"])
        if not s:
            sorun.append(f"{y['id']}: sınıflandırılmadı")
            continue
        konular = {k["konu"] for k in s["konular"]}
        if not s["konular"]:
            sorun.append(f"{y['id']}: konu yok")
        for k in s["konular"]:
            if not metinde_var(k["alinti"], y["yorum"]):
                sorun.append(f"{y['id']}: '{k['konu']}' alıntısı yorumda yok: '{k['alinti'][:50]}'")
        for i in isaretler(y["yorum"]):
            if i not in konular:
                sorun.append(f"{y['id']}: yorumda '{i}' ifadesi var ('{y['yorum'][:60]}') ama bu konu işaretlenmemiş")
    return sorun


def uret(yorum: list[dict], duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    sonuc = {}
    hedef = yorum
    if duzeltme is not None:  # düzeltmede yalnız sorunlu yorumlar yeniden sorulur
        sorunlu = {s.split(":")[0] for s in duzeltme}
        sonuc.update({a: b for a, b in (onceki or {}).items() if a not in sorunlu})
        hedef = [y for y in yorum if y["id"] in sorunlu]
    for j in range(0, len(hedef), PAKET):
        parca = hedef[j:j + PAKET]
        ids = {y["id"] for y in parca}
        satir = "\n".join(f'<yorum id="{y["id"]}" urun="{y["ad"] or y["urun"]}" puan="{puan_yaz(y)}" isaretler="{", ".join(isaretler(y["yorum"])) or "-"}">'
                          f"{llm.maskele(y['yorum'])}</yorum>" for y in parca)
        d = [s for s in (duzeltme or []) if s.split(":")[0] in ids]
        mesaj = f"<konular>{'; '.join(KONULAR)}</konular>\n<yorumlar>\n{satir}\n</yorumlar>" + duzeltme_blogu(d, [b for a, b in (onceki or {}).items() if a in ids])
        for s in llm.json_iste(sistem, mesaj, SEMA).get("yorumlar", []):
            if s.get("id") in ids:
                sonuc[s["id"]] = s
    return sonuc


def tutarsiz(y: dict, s: dict) -> bool:
    return y["puan"] is not None and ((y["puan"] >= 4 and s["duygu"] == "Olumsuz") or (y["puan"] <= 2 and s["duygu"] == "Olumlu"))


def urun_ozeti(yorum: list[dict], sonuc: dict) -> list[dict]:
    gr = defaultdict(list)
    for y in yorum:
        if y["id"] in sonuc:
            gr[y["urun"]].append(y)
    out = []
    for u, l in gr.items():
        puanlar = [y["puan"] for y in l if y["puan"] is not None]
        konu = defaultdict(lambda: {"Olumlu": 0, "Nötr": 0, "Olumsuz": 0, "alinti": []})
        for y in l:
            for k in sonuc[y["id"]]["konular"]:
                konu[k["konu"]][k["duygu"]] += 1
                konu[k["konu"]]["alinti"].append((k["duygu"], f"“{k['alinti']}” ({y['id']})"))
        en_kotu = max(konu.items(), key=lambda x: (x[1]["Olumsuz"], -x[1]["Olumlu"]), default=(None, None))
        out.append({"urun": u, "ad": l[0]["ad"], "yorum": len(l), "ort": round(sum(puanlar) / len(puanlar), 2) if puanlar else None, "konu": dict(konu),
                    "sikayet": en_kotu[0] if en_kotu[1] and en_kotu[1]["Olumsuz"] else "",
                    "guvenlik": [y["id"] for y in l if any(k["konu"] == GUVENLIK_K for k in sonuc[y["id"]]["konular"])]})
    return sorted(out, key=lambda r: (-len(r["guvenlik"]), r["ort"] if r["ort"] is not None else 9))


def rapor_yaz(cikti: Path, yorum, sonuc, ozet, kalan, tur) -> None:
    wb = Workbook()
    a = wb.active
    a.title = "Acil"
    baslik(a, ["Yorum No", "Tarih", "Ürün", "Puan", "Yorum", "Yapılacak"], (9, 11, 20, 6, 70, 60))
    for y in yorum:
        s = sonuc.get(y["id"])
        if s and any(k["konu"] == GUVENLIK_K for k in s["konular"]):
            a.append([y["id"], y["tarih"], y["ad"] or y["urun"], y["puan"], y["yorum"],
                      "Müşteriyle hemen iletişime geçin; ürünü / partiyi kalite ve ürün güvenliği birimine bildirin; benzer şikâyetleri tarayın"])
            a.cell(a.max_row, 2).number_format = "DD.MM.YYYY"
            boya(a, [1], KIRMIZI)
    if a.max_row == 1:
        a.append(["—", "", "", "", "Ürün güvenliği belirtisi olan yorum yok"])
    hizala(a)
    o = wb.create_sheet("Ürün Özeti")
    baslik(o, ["Ürün Kodu", "Ürün Adı", "Yorum", "Ort. Puan", "En Çok Şikâyet", "Konu", "Olumlu", "Olumsuz", "Net", "Örnek Alıntılar"], (10, 22, 7, 9, 22, 26, 8, 8, 6, 90))
    for r in ozet:
        ilk = True
        for k, v in sorted(r["konu"].items(), key=lambda x: -(x[1]["Olumlu"] + x[1]["Olumsuz"])):
            o.append(([r["urun"], r["ad"], r["yorum"], r["ort"], r["sikayet"]] if ilk else [""] * 5)
                     + [k, v["Olumlu"], v["Olumsuz"], v["Olumlu"] - v["Olumsuz"], " · ".join(x for _, x in v["alinti"][:3])])
            boya(o, [9], YESIL if v["Olumlu"] > v["Olumsuz"] else (KIRMIZI if v["Olumsuz"] > v["Olumlu"] else SARI))
            if k == GUVENLIK_K:
                boya(o, [6], KIRMIZI)
            ilk = False
    son_not(o, "Önemli: Duygu ve konu etiketleri yapay zekâ sınıflandırmasıdır. Ürün güvenliği yorumlarını mutlaka okuyun.", kalan, tur)
    hizala(o)
    y2 = wb.create_sheet("Yorumlar")
    baslik(y2, ["Yorum No", "Tarih", "Ürün", "Puan", "Yorum", "Duygu", "Konular", "Özet", "Puan-Duygu Çelişkisi"], (9, 11, 20, 6, 60, 9, 44, 40, 14))
    for y in yorum:
        s = sonuc.get(y["id"])
        if not s:
            y2.append([y["id"], y["tarih"], y["ad"] or y["urun"], y["puan"], y["yorum"], "AI yanıt vermedi"])
            continue
        y2.append([y["id"], y["tarih"], y["ad"] or y["urun"], y["puan"], y["yorum"], s["duygu"], "; ".join(f"{k['konu']} ({k['duygu']})" for k in s["konular"]),
                   s["ozet"], "Kontrol edin" if tutarsiz(y, s) else ""])
        y2.cell(y2.max_row, 2).number_format = "DD.MM.YYYY"
        boya(y2, [6, 7, 8], AI_DOLGU)
        if tutarsiz(y, s):
            boya(y2, [9], SARI)
    y2.freeze_panes = "B2"
    y2.auto_filter.ref = f"A1:I{y2.max_row}"
    hizala(y2)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(girdi: Path, cikti: Path, tur: int = 2, evet: bool = False) -> dict:
    yorum = yorumlari_oku(girdi)
    if not yorum:
        raise llm.LLMHatasi(f"{girdi.name}: yorum yok.")
    print(f"[OK] {len(yorum)} yorum · güvenlik belirtisi {sum(GUVENLIK_K in isaretler(y['yorum']) for y in yorum)}")
    llm.onay_al(f"{len(yorum)} yorum metni, ürün adı ve puanı gönderilecek (telefon / e-posta maskeli).", evet)
    sonuc, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(yorum, d, o), lambda x: denetle(yorum, x), tur)
    ozet = urun_ozeti(yorum, sonuc)
    rapor_yaz(cikti, yorum, sonuc, ozet, kalan, yapilan)
    return {"yorumlar": yorum, "sonuc": sonuc, "ozet": ozet, "kalan": kalan, "tur": yapilan}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    p = argparse.ArgumentParser(description="Ürün yorumlarını duygu ve konu bazında sınıflandırıp ürün bazında özetler.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "yorumlar.csv",
                   help="Yorumlar (.xlsx/.csv): Yorum; Yorum No, Tarih, Ürün Kodu, Ürün Adı, Puan isteğe bağlı")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "urun_degerlendirme.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.cikti, a.tur, a.evet)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    for r in s["ozet"]:
        print(f"[OK] {r['urun']} {r['ad']}: {r['yorum']} yorum · ort. {r['ort']} · en çok şikâyet: {r['sikayet'] or '-'}"
              + (f" · ÜRÜN GÜVENLİĞİ: {', '.join(r['guvenlik'])}" if r["guvenlik"] else ""))
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
