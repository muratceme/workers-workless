"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri" / "yorumlar.csv"
MESAJLAR = []
KAL, BED, KAR, PAK, FIY, ACI, SAT, GUV = agent.KONULAR[:8]
P, N, O = "Olumlu", "Nötr", "Olumsuz"
DOGRU = {
    "Y-1": (P, [(KAL, P, "Çok hızlı kaynatıyor"), (KAR, P, "Kargo da ertesi gün geldi")]),
    "Y-2": (O, [(GUV, O, "fişten duman çıktı")]),
    "Y-3": (P, [(KAL, N, "kapağı biraz zor açılıyor")]),
    "Y-4": (O, [(PAK, O, "Kutusu ezik geldi"), (KAL, O, "ürün çizik")]),
    "Y-5": (P, [(FIY, P, "Fiyatına göre çok iyi")]),
    "Y-6": (P, [(KAL, P, "Çok yumuşak"), (ACI, P, "renkleri fotoğraftaki gibi")]),
    "Y-7": (O, [(KAL, O, "çok tüy döktü, rengi soldu")]),
    "Y-8": (N, [(KAL, N, "idare eder"), (KAR, O, "kargo 9 günde geldi")]),
    "Y-9": (P, [(ACI, O, "açıklamada yazan ölçüden küçük çıktı")]),
    "Y-10": (P, [(BED, O, "Kalıbı dar"), (KAL, P, "Kumaşı güzel")]),
    "Y-11": (O, [(GUV, O, "cildinde kızarıklık ve kaşıntı oldu")]),
    "Y-12": (P, [(SAT, P, "sorularıma hemen cevap verdi")]),
    "Y-13": (O, [(BED, O, "Dar geldi")]),
}


def sonuc(dogru=True) -> dict:
    d = {i: {"id": i, "duygu": g, "ozet": "özet", "konular": [{"konu": k, "duygu": du, "alinti": a} for k, du, a in ks]} for i, (g, ks) in DOGRU.items()}
    if not dogru:
        d["Y-2"]["konular"] = [{"konu": KAL, "duygu": O, "alinti": "fişten duman çıktı"}]
        d["Y-11"]["konular"] = [{"konu": KAL, "duygu": O, "alinti": "Baskısından"}]
        d["Y-8"]["konular"][1]["alinti"] = "kargo 10 günde geldi"
    return d


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    d = sonuc("<duzeltme>" in kullanici)
    return {"yorumlar": [d[i] for i in re.findall(r'<yorum id="([^"]+)"', kullanici)]}


class KodTesti(unittest.TestCase):
    def test_isaretler(self):
        y = agent.yorumlari_oku(ORNEK)
        self.assertEqual({x["id"]: agent.isaretler(x["yorum"]) for x in y if agent.isaretler(x["yorum"])},
                         {"Y-1": [KAR], "Y-2": [GUV], "Y-4": [PAK], "Y-8": [KAR], "Y-11": [GUV]})

    def test_denetim_ve_ozet(self):
        y = agent.yorumlari_oku(ORNEK)
        self.assertEqual(agent.denetle(y, sonuc()), [])
        s = agent.denetle(y, sonuc(False))
        self.assertEqual(len(s), 3)
        self.assertTrue(any("Y-2" in x and GUV in x for x in s) and any("Y-11" in x for x in s) and any("10 günde" in x for x in s))
        o = agent.urun_ozeti(y, sonuc())
        self.assertEqual([(r["urun"], r["ort"], r["guvenlik"]) for r in o], [("TSR-03", 3.0, ["Y-11"]), ("KTL-01", 3.4, ["Y-2"]), ("HVL-02", 3.75, [])])
        self.assertEqual((o[0]["sikayet"], o[2]["konu"][KAL]["Olumsuz"]), (BED, 1))
        self.assertEqual([x["id"] for x in y if agent.tutarsiz(x, sonuc()[x["id"]])], [])
        d = sonuc()
        d["Y-9"]["duygu"] = O
        self.assertTrue(agent.tutarsiz(y[8], d["Y-9"]))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "u.xlsx"
            s = agent.calistir(ORNEK, cikti, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            self.assertIn('puan="1" isaretler="Ürün güvenliği"', MESAJLAR[0])
            self.assertEqual(re.findall(r'<yorum id="([^"]+)"', MESAJLAR[1]), ["Y-2", "Y-8", "Y-11"])
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Acil", "Ürün Özeti", "Yorumlar"])
            self.assertEqual([wb["Acil"].cell(r, 1).value for r in (2, 3)], ["Y-2", "Y-11"])

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(ORNEK, Path(tmp) / "u.xlsx", tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 3)
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--girdi", str(ORNEK.with_name("yok.csv")), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
