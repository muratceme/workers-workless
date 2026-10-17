"""
Öğrenci Gelişim Raporu Taslağı — Gerisi Bende AI Agent
Akademik › Rehber Öğretmen

1. Kod not tablosunu (Öğrenci No; Ad Soyad; Sınıf; Ders; Önceki Dönem; Bu Dönem), devamsızlığı ve öğretmen gözlemlerini
   ("1021: …" satırları) okur ve öğrenci başına hesaplar: ders puanları ve değişimleri, bu dönem ortalaması, belirgin
   düşüşler (--dusus puan ve üzeri). Öğrenci adı ve numarası modele gönderilmez ({OGRENCI}, Ö1, Ö2 …).
2. Model her öğrenci için güçlü yönler, gelişime açık alanlar, öneriler, veli görüşmesinde konuşulacak sorular ve genel
   değerlendirme yazar; her madde dayandığı ders ya da gözlem numarasını gösterir.
3. Kod denetler:
   - Dayanak ders öğrencinin derslerinden mi, gözlem numarası geçerli mi
   - Metindeki sayılar öğrencinin puan, değişim, ortalama ve devamsızlık değerlerinden mi
   - Güçlü yön olarak gösterilen ders ortalamanın altında mı; belirgin düşüş gösteren her ders gelişime açık alanlarda var mı
   - Tanı koyan ya da etiketleyen ifade ("dikkat eksikliği", "hiperaktif", "tembel") ve başka öğrencilerle kıyas var mı
   Sorun varsa yalnız o öğrencinin taslağı modele geri gönderilir (en fazla --tur kez).

Kullanım:
    python agent.py                                                      # örnek sınıf
    python agent.py --notlar notlar.xlsx --devamsizlik devamsizlik.xlsx --gozlemler gozlemler.txt --dusus 10
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
_M = {"type": "object", "properties": {"metin": S, "dayanak": {"type": "array", "items": S}}, "required": ["metin", "dayanak"], "additionalProperties": False}
SEMA = {
    "type": "object",
    "properties": {"genel_degerlendirme": S, "guclu_yonler": {"type": "array", "items": _M}, "gelisime_acik": {"type": "array", "items": _M},
                   "oneriler": {"type": "array", "items": _M}, "veli_sorulari": {"type": "array", "items": S}},
    "required": ["genel_degerlendirme", "guclu_yonler", "gelisime_acik", "oneriler", "veli_sorulari"],
    "additionalProperties": False,
}
ETIKET = re.compile(r"\b(dikkat eksikli\w*|hiperaktiv\w*|hiperaktif\w*|dehb|disleksi\w*|diskalkuli\w*|otizm\w*|otistik|depresyon\w*|anksiyete bozuklu\w*|"
                    r"ogrenme guclugu (?:var|vardir)|tembel\w*|yeteneksiz\w*|zekasi (?:dusuk|yetersiz)|aptal\w*|sorunlu ogrenci|basarisiz ogrenci)\b")
KIYAS = re.compile(r"\b(sinifin en\w*|siniftaki en\w*|diger ogrenciler\w*|arkadaslarindan (?:daha|geri)\w*|arkadaslarina gore|akranlarina gore)\b")


def puan(v) -> float | None:
    try:
        return float(v) if isinstance(v, (int, float)) else float(sayi_norm(str(v))) if str(v).strip() else None
    except ValueError:
        return None


def tr(v: float) -> str:
    return sayi_yaz(round(v, 1)).replace(".", ",")


def ogrencileri_oku(not_yolu: Path, dev_yolu: Path | None, gozlem_yolu: Path | None, dusus: float) -> list[dict]:
    r = kayitlar(not_yolu, {"no": ("Öğrenci No", "Okul No", "No"), "ad": ("Ad Soyad", "Öğrenci", "Öğrenci Adı"), "sinif": ("Sınıf", "Şube"), "ders": ("Ders",),
                            "onceki": ("Önceki Dönem", "Önceki"), "bu": ("Bu Dönem", "Güncel", "Puan")}, ("no", "ders", "bu"))
    ogr = {}
    for x in r:
        no = str(x["no"]).strip()
        if not no or puan(x["bu"]) is None:
            continue
        o = ogr.setdefault(no, {"no": no, "ad": str(x["ad"]).strip(), "sinif": str(x["sinif"]).strip(), "dersler": [], "ozursuz": None, "ozurlu": None,
                                "gozlemler": []})
        o["dersler"].append({"ders": str(x["ders"]).strip(), "onceki": puan(x["onceki"]), "bu": puan(x["bu"])})
    if dev_yolu:
        for x in kayitlar(dev_yolu, {"no": ("Öğrenci No", "Okul No", "No"), "ozursuz": ("Özürsüz", "Özürsüz Gün"), "ozurlu": ("Özürlü", "Özürlü Gün")}, ("no",)):
            o = ogr.get(str(x["no"]).strip())
            if o:
                o["ozursuz"], o["ozurlu"] = puan(x["ozursuz"]), puan(x["ozurlu"])
    if gozlem_yolu:
        if not gozlem_yolu.exists():
            raise llm.LLMHatasi(f"{gozlem_yolu}: gözlem dosyası bulunamadı.")
        for s in belge.metin_oku(gozlem_yolu).splitlines():
            m = re.match(r"^\s*(\w+)\s*[:\-–]\s*(.+)$", s)
            if m and m.group(1) in ogr:
                ogr[m.group(1)]["gozlemler"].append(m.group(2).strip())
    out = []
    for i, o in enumerate(ogr.values(), 1):
        o["kimlik"] = f"Ö{i}"
        o["ortalama"] = sum(d["bu"] for d in o["dersler"]) / len(o["dersler"])
        for d in o["dersler"]:
            d["fark"] = None if d["onceki"] is None else d["bu"] - d["onceki"]
        o["dusenler"] = [d["ders"] for d in o["dersler"] if d["fark"] is not None and d["fark"] <= -dusus]
        ilk = o["ad"].split()[0] if o["ad"] else ""
        o["gozlemler"] = [llm.maskele(re.sub(rf"\b{re.escape(ilk)}\b", "{OGRENCI}", g) if ilk else g) for g in o["gozlemler"]]
        out.append(o)
    return out


def rakamlar(o: dict) -> set[str]:
    s = {tr(o["ortalama"]), str(len(o["dersler"]))}
    for d in o["dersler"]:
        s |= {tr(d["bu"])} | ({tr(d["onceki"]), tr(abs(d["fark"]))} if d["fark"] is not None else set())
    for a in ("ozursuz", "ozurlu"):
        if o[a] is not None:
            s.add(tr(o[a]))
    s = set().union(*(sayilar(x) for x in s))
    return s.union(*(sayilar(g) for g in o["gozlemler"]))


def denetle(ogrenciler: list[dict], sonuc: dict) -> list[str]:
    sorun = []
    for o in ogrenciler:
        a, t = o["kimlik"], sonuc.get(o["kimlik"])
        if not t:
            sorun.append(f"{a}: taslak yazılmadı")
            continue
        dersler = {katla(d["ders"]): d for d in o["dersler"]}
        izinli = rakamlar(o)
        gruplar = [("Güçlü yön", t["guclu_yonler"]), ("Gelişime açık", t["gelisime_acik"]), ("Öneri", t["oneriler"])]
        for ad, liste in gruplar:
            for i, m in enumerate(liste, 1):
                yer = f"{ad} {i}"
                if not m["dayanak"]:
                    sorun.append(f"{a}: {yer}: dayanak (ders ya da gözlem) yok")
                for d in m["dayanak"]:
                    g = re.fullmatch(r"G(\d+)", d.strip())
                    if g and not 1 <= int(g.group(1)) <= len(o["gozlemler"]):
                        sorun.append(f"{a}: {yer}: olmayan gözlem {d}")
                    elif not g and katla(d) not in dersler:
                        sorun.append(f"{a}: {yer}: öğrencinin dersi değil: '{d}'")
                    elif not g and ad == "Güçlü yön" and dersler[katla(d)]["bu"] < o["ortalama"]:
                        sorun.append(f"{a}: {yer}: {d} ({tr(dersler[katla(d)]['bu'])}) ortalamanın ({tr(o['ortalama'])}) altında; güçlü yön gösterilemez")
        metinler = [("Genel", t["genel_degerlendirme"])] + [(ad, m["metin"]) for ad, liste in gruplar for m in liste] + [("Soru", s) for s in t["veli_sorulari"]]
        for ad, m in metinler:
            fazla = sorted(sayilar(m) - izinli)
            if fazla:
                sorun.append(f"{a}: {ad}: öğrencinin verilerinde olmayan sayı: {', '.join(fazla)}")
            if ETIKET.search(katla(m)):
                sorun.append(f"{a}: {ad}: tanı koyan / etiketleyen ifade ('{ETIKET.search(katla(m)).group(0)}'); gözlem diliyle yazın")
            if KIYAS.search(katla(m)):
                sorun.append(f"{a}: {ad}: başka öğrencilerle kıyas ('{KIYAS.search(katla(m)).group(0)}')")
        acik = {katla(d) for m in t["gelisime_acik"] for d in m["dayanak"]}
        eksik = [d for d in o["dusenler"] if katla(d) not in acik]
        if eksik:
            sorun.append(f"{a}: belirgin düşüş gösteren ders gelişime açık alanlarda yok: {', '.join(eksik)}")
        if "{OGRENCI}" not in t["genel_degerlendirme"]:
            sorun.append(f"{a}: genel değerlendirmede {{OGRENCI}} yok")
    return sorun


def ogrenci_metni(o: dict) -> str:
    d = "\n".join(f"- {x['ders']}: önceki {tr(x['onceki']) if x['onceki'] is not None else '-'} → bu dönem {tr(x['bu'])}"
                  + (f" (değişim {'+' if x['fark'] > 0 else '−' if x['fark'] < 0 else '±'}{tr(abs(x['fark']))})" if x["fark"] is not None else "") for x in o["dersler"])
    g = "\n".join(f"G{i}. {s}" for i, s in enumerate(o["gozlemler"], 1)) or "(gözlem yok)"
    dev = f"özürsüz {tr(o['ozursuz'])} gün, özürlü {tr(o['ozurlu'] or 0)} gün" if o["ozursuz"] is not None else "bilgi yok"
    return (f"<ogrenci kimlik=\"{o['kimlik']}\" ad=\"{{OGRENCI}}\">\nBu dönem ortalaması: {tr(o['ortalama'])}\nBelirgin düşüş: {', '.join(o['dusenler']) or 'yok'}\n"
            f"Devamsızlık: {dev}\nDersler:\n{d}\nÖğretmen gözlemleri:\n{g}\n</ogrenci>")


def uret(ogrenciler: list[dict], duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    sonuc, hedef = {}, ogrenciler
    if duzeltme is not None:  # düzeltmede yalnız sorunlu öğrencilerin taslağı yeniden yazılır
        sorunlu = {s.split(":")[0] for s in duzeltme}
        sonuc.update({a: b for a, b in (onceki or {}).items() if a not in sorunlu})
        hedef = [o for o in ogrenciler if o["kimlik"] in sorunlu]
    for o in hedef:
        d = [s for s in (duzeltme or []) if s.split(":")[0] == o["kimlik"]]
        sonuc[o["kimlik"]] = llm.json_iste(sistem, ogrenci_metni(o) + duzeltme_blogu(d, (onceki or {}).get(o["kimlik"])), SEMA)
    return sonuc


def rapor_metni(o: dict, t: dict) -> str:
    ad = o["ad"] or o["no"]
    f = lambda s: s.replace("{OGRENCI}", ad)  # noqa: E731
    s = [f"ÖĞRENCİ GELİŞİM RAPORU (TASLAK) — {ad} · {o['sinif']} · No: {o['no']}", "", f(t["genel_degerlendirme"]), "", "Güçlü yönler:"]
    s += [f"- {f(m['metin'])}" for m in t["guclu_yonler"]] + ["", "Gelişime açık alanlar:"] + [f"- {f(m['metin'])}" for m in t["gelisime_acik"]]
    s += ["", "Öneriler:"] + [f"- {f(m['metin'])}" for m in t["oneriler"]] + ["", "Veli görüşmesinde konuşulacaklar:"] + [f"- {f(x)}" for x in t["veli_sorulari"]]
    return "\n".join(s + ["", "Bu taslak rehber öğretmen tarafından gözden geçirilmeden veliyle paylaşılmaz."])


def rapor_yaz(cikti: Path, ogrenciler, sonuc, kalan, tur) -> list[Path]:
    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    baslik(ws, ["Öğrenci No", "Ad Soyad", "Sınıf", "Ortalama", "Belirgin Düşüş", "Özürsüz", "Özürlü", "Güçlü Yönler", "Gelişime Açık", "Kontrol"],
           (10, 20, 7, 9, 22, 8, 8, 50, 50, 30))
    dosyalar = []
    for o in ogrenciler:
        t = sonuc.get(o["kimlik"])
        sor = [s.split(": ", 1)[1] for s in kalan if s.startswith(o["kimlik"] + ":")]
        f = lambda s: s.replace("{OGRENCI}", o["ad"])  # noqa: E731
        ws.append([o["no"], o["ad"], o["sinif"], round(o["ortalama"], 1), ", ".join(o["dusenler"]), o["ozursuz"], o["ozurlu"],
                   "\n".join(f(m["metin"]) for m in t["guclu_yonler"]) if t else "", "\n".join(f(m["metin"]) for m in t["gelisime_acik"]) if t else "",
                   "; ".join(sor) or ("Geçti" if t else "AI yanıt vermedi")])
        boya(ws, [8, 9], AI_DOLGU)
        boya(ws, [10], KIRMIZI if sor or not t else YESIL)
        if o["dusenler"]:
            boya(ws, [5], SARI)
        if t and not sor:
            yol = cikti.with_name(f"{cikti.stem}_{re.sub(r'[^A-Za-z0-9]+', '_', o['no'])}.txt")
            yol.parent.mkdir(parents=True, exist_ok=True)
            yol.write_text(rapor_metni(o, t), encoding="utf-8")
            dosyalar.append(yol)
    son_not(ws, "Önemli: Taslaklar rehber öğretmen tarafından gözden geçirilmeli; öğrenci hakkında tanı konulmaz, gerekirse ilgili uzmana "
                "yönlendirme okul idaresi ve veliyle birlikte değerlendirilir.", kalan, tur)
    hizala(ws)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    return dosyalar


def calistir(not_yolu: Path, cikti: Path, dev_yolu: Path | None = None, gozlem_yolu: Path | None = None, dusus: float = 10, tur: int = 2,
             evet: bool = False) -> dict:
    ogrenciler = ogrencileri_oku(not_yolu, dev_yolu, gozlem_yolu, dusus)
    if not ogrenciler:
        raise llm.LLMHatasi(f"{not_yolu.name}: öğrenci yok.")
    print(f"[OK] {len(ogrenciler)} öğrenci · {sum(len(o['gozlemler']) for o in ogrenciler)} gözlem · belirgin düşüş: "
          f"{sum(len(o['dusenler']) for o in ogrenciler)} ders")
    llm.onay_al(f"{len(ogrenciler)} öğrencinin ders puanları, devamsızlığı ve öğretmen gözlemleri gönderilecek (ad ve okul numarası gönderilmez).", evet)
    sonuc, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(ogrenciler, d, o), lambda x: denetle(ogrenciler, x), tur)
    dosyalar = rapor_yaz(cikti, ogrenciler, sonuc, kalan, yapilan)
    return {"ogrenciler": ogrenciler, "sonuc": sonuc, "kalan": kalan, "tur": yapilan, "dosyalar": dosyalar}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    o = BURASI / "ornek_veri"
    p = argparse.ArgumentParser(description="Not, devamsızlık ve öğretmen gözlemlerinden veli görüşmesi için öğrenci gelişim raporu taslağı hazırlar.")
    p.add_argument("--notlar", type=Path, default=o / "notlar.csv", help="Not tablosu (.xlsx/.csv): Öğrenci No, Ders, Bu Dönem; Ad Soyad, Sınıf, Önceki Dönem")
    p.add_argument("--devamsizlik", type=Path, default=o / "devamsizlik.csv", help="Devamsızlık (.xlsx/.csv): Öğrenci No; Özürsüz, Özürlü (isteğe bağlı)")
    p.add_argument("--gozlemler", type=Path, default=o / "gozlemler.txt", help="Öğretmen gözlemleri (.txt/.docx): '1021: gözlem' satırları (isteğe bağlı)")
    p.add_argument("--dusus", type=float, default=10, help="Bu kadar puan ve üzeri düşüş belirgin sayılır (varsayılan 10)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "gelisim_raporlari.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.notlar, a.cikti, a.devamsizlik, a.gozlemler, a.dusus, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(s['dosyalar'])}/{len(s['ogrenciler'])} öğrenci için taslak hazır")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
