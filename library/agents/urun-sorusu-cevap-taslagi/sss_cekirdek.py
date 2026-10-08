"""
Sık Sorulan Sorulara Cevap Taslağı — Workers / Workless AI Agent
Bilgi Teknolojileri › BT Destek Uzmanı

1. Bilgi bankası (tek dosya veya .md/.txt/.docx/.pdf klasörü) başlıklarına göre bölümlere ayrılır.
2. Her soru için kod, en ilgili bölümleri BM25 aramasıyla bulur (Türkçe için 5 harfli kök kesme). Modele
   bilgi bankasının tamamı değil, yalnız bu bölümler gönderilir. Sorudaki e-posta, telefon, IBAN, TCKN ve
   "şifre: ..." gibi parolalar maskelenir.
3. Model YALNIZ verilen bölümlere dayanarak cevap taslağı yazar, kullandığı bölümleri [K3] gibi belirtir;
   bölümlerde cevap yoksa bunu söyler.
4. Kod denetler: kaynak göstermeyen cevap "kapsam yok" sayılır; kullanıcıdan parola isteyen taslak reddedilir;
   güvenlik olayı belirtisi taşıyan sorular her zaman insana yönlendirilir; arıza gibi görünen sorular için
   talep açılması önerilir. Cevaplanamayan sorular "Bilgi Bankası Boşlukları" sayfasında toplanır.
5. Sonuç Excel'e yazılır; her taslak için 'Onay' sütunu vardır. Hiçbir cevap otomatik gönderilmez.

Kullanım:
    python agent.py                                              # örnek sorularla dener
    python agent.py --girdi sorular.xlsx --bilgi ./bilgi_bankasi
    python agent.py --soru "VPN'e evden nasıl bağlanırım?" --bilgi ./bilgi_bankasi
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
PAKET = 10                       # bir istekte en fazla soru
SORU_SINIRI = 2000               # tek soruda gönderilen en fazla karakter
BOLUM_SINIRI = 1800              # başlıksız metinlerde bölüm boyu (karakter)
KOK = 5                          # Türkçe için kök kesme uzunluğu

DURAK = {"ve", "ile", "bir", "bu", "şu", "da", "de", "mi", "mı", "mu", "mü", "ne", "nasıl", "için", "gibi", "ama", "veya",
         "ya", "çok", "daha", "en", "her", "hiç", "ben", "biz", "siz", "sen", "o", "benim", "bana", "beni", "var", "yok",
         "olarak", "olan", "ise", "kadar", "sonra", "önce", "acaba", "merhaba", "teşekkürler", "rica", "lütfen", "selam",
         "iyi", "günler", "çalışmalar", "hocam", "arkadaşlar"}

GUVENLIK = re.compile(r"(oltalama|phishing|şüpheli (?:e-?posta|mail|bağlantı|link)|virüs|fidye|ransomware|hacklendi|"
                      r"hesabım(?:a)? (?:ele geçirildi|başkası)|bağlantıya tıkladım|linke tıkladım|şifremi (?:girdim|kaptırdım))", re.I)
ARIZA = re.compile(r"(çalışmıyor|açılmıyor|hata veriyor|hata alıyorum|bozuldu|erişemiyorum|bağlanamıyorum|donuyor|"
                   r"kilitlendi|yanıt vermiyor|çöktü)", re.I)
# "şifre: X", "parolam = X", "şifrem olarak X" ya da "şifrem Ankara2026" (harf + rakam içeren sözcük)
PAROLA = re.compile(r"((?:şifre|parola|password|pin)(?:m|miz)?\s*(?:[:=]|olarak)\s*)(\S{4,})|"
                    r"((?:şifrem|parolam|şifremiz|parolamız)\s+)((?=\S*\d)(?=\S*[^\d\s])\S{4,})", re.I)
PAROLA_ISTEME = re.compile(r"(şifrenizi|parolanızı|mevcut şifre|eski şifre|pin kodunuzu)[^.\n]{0,40}"
                           r"(yazın|yazınız|gönderin|gönderiniz|iletin|iletiniz|paylaşın|paylaşınız|bildirin|bildiriniz|söyleyin)", re.I)


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def katla(s) -> str:
    s = kucuk(s)
    for a, b in zip("ıöüşğçâîû", "iousgcaiu"):
        s = s.replace(a, b)
    return s


DURAK_K = {katla(d) for d in DURAK}


def kokler(metin: str) -> list[str]:
    """Türkçe için basit ve sağlam bir arama yaklaşımı: kelimeyi ilk 5 harfine kes (F5)."""
    return [w[:KOK] for w in re.findall(r"[a-z0-9]+", katla(metin)) if len(w) > 1 and w not in DURAK_K]


@dataclass
class Bolum:
    id: str
    kaynak: str
    baslik: str
    metin: str
    kok: list[str] = field(default_factory=list)


@dataclass
class Soru:
    id: str
    tarih: datetime | None
    kullanici: str
    departman: str
    metin: str
    etiketler: list[str] = field(default_factory=list)
    eslesme: list[tuple[str, float]] = field(default_factory=list)


# ----------------------------------------------------------------------------
# Bilgi bankası
# ----------------------------------------------------------------------------

def bolumle(kaynak: str, metin: str) -> list[tuple[str, str]]:
    """Markdown başlıklarına (#, ##, ###) göre; başlık yoksa paragrafları birleştirerek bölümler."""
    satirlar = metin.replace("\r\n", "\n").split("\n")
    if any(re.match(r"^#{1,4}\s+\S", s) for s in satirlar):
        bolumler, yol, govde = [], [], []

        def kapat():
            icerik = "\n".join(govde).strip()
            if icerik:
                bolumler.append((" › ".join(yol) or kaynak, icerik))
            govde.clear()

        for s in satirlar:
            m = re.match(r"^(#{1,4})\s+(.+?)\s*#*\s*$", s)
            if m:
                kapat()
                seviye = len(m.group(1))
                yol = yol[:seviye - 1] + [m.group(2).strip()]
            else:
                govde.append(s)
        kapat()
        return bolumler
    paragraflar = [p.strip() for p in re.split(r"\n\s*\n", metin) if p.strip()]
    bolumler, parca = [], ""
    for p in paragraflar:
        if parca and len(parca) + len(p) > BOLUM_SINIRI:
            bolumler.append(parca)
            parca = ""
        parca = f"{parca}\n\n{p}" if parca else p
    if parca:
        bolumler.append(parca)
    return [(f"{kaynak} · bölüm {i}", b) for i, b in enumerate(bolumler, 1)]


def bilgi_bankasi_oku(yol: Path) -> list[Bolum]:
    dosyalar = sorted(p for p in yol.rglob("*") if p.suffix.lower() in belge.DESTEKLENEN) if yol.is_dir() else [yol]
    if not dosyalar:
        raise llm.LLMHatasi(f"{yol}: bilgi bankasında okunacak .md/.txt/.docx/.pdf dosyası yok.")
    bolumler = []
    for d in dosyalar:
        try:
            metin = belge.metin_oku(d)
        except belge.BelgeHatasi as h:
            print(f"[!] {h}")
            continue
        for baslik, icerik in bolumle(d.name, metin):
            b = Bolum(f"K{len(bolumler) + 1}", d.name, baslik, icerik)
            b.kok = kokler(f"{baslik}\n{baslik}\n{icerik}")        # başlık iki kez: başlık eşleşmesi daha değerli
            bolumler.append(b)
    if not bolumler:
        raise llm.LLMHatasi(f"{yol}: bilgi bankasından metin çıkarılamadı.")
    return bolumler


class BM25:
    def __init__(self, belgeler: list[list[str]], k1: float = 1.2, b: float = 0.75):
        self.k1, self.b = k1, b
        self.tf = [Counter(d) for d in belgeler]
        self.uz = [len(d) for d in belgeler]
        self.ort = sum(self.uz) / len(self.uz) if self.uz else 1
        n = len(belgeler)
        df = Counter(t for d in belgeler for t in set(d))
        self.idf = {t: math.log(1 + (n - v + 0.5) / (v + 0.5)) for t, v in df.items()}

    def puan(self, sorgu: list[str]) -> list[float]:
        sonuc = []
        for tf, uz in zip(self.tf, self.uz):
            p = 0.0
            for t in set(sorgu):
                f = tf.get(t, 0)
                if f:
                    p += self.idf[t] * f * (self.k1 + 1) / (f + self.k1 * (1 - self.b + self.b * uz / self.ort))
            sonuc.append(p)
        return sonuc


def eslestir(sorular: list[Soru], bolumler: list[Bolum], k: int, esik: float) -> None:
    motor = BM25([b.kok for b in bolumler])
    for s in sorular:
        puanlar = motor.puan(kokler(s.metin))
        sira = sorted(range(len(bolumler)), key=lambda i: -puanlar[i])
        s.eslesme = [(bolumler[i].id, puanlar[i]) for i in sira[:k] if puanlar[i] >= esik]


# ----------------------------------------------------------------------------
# Sorular
# ----------------------------------------------------------------------------

ALANLAR = {
    "id": ("soru no", "talep no", "no", "id", "ıd", "ticket"),
    "tarih": ("tarih", "geliş tarihi"),
    "kullanici": ("kullanıcı", "gönderen", "çalışan", "ad soyad"),
    "departman": ("departman", "birim", "bölüm"),
    "metin": ("soru", "mesaj", "metin", "açıklama", "içerik"),
}


def tarih_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(x).strip(), f)
        except ValueError:
            pass
    return None


def sorulari_oku(yol: Path) -> list[Soru]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = belge.txt_oku(yol)
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    s = [r for r in s if any(c not in (None, "") for c in r)]
    if not s:
        raise llm.LLMHatasi(f"{yol}: dosya boş.")
    b = [katla(x) for x in s[0]]
    k = {alan: next((i for i, x in enumerate(b) if x in {katla(a) for a in adlar}), None) for alan, adlar in ALANLAR.items()}
    if k["metin"] is None:
        raise llm.LLMHatasi(f"Soru dosyasında 'Soru' (veya 'Mesaj') sütunu bulunamadı. Başlıklar: {s[0]}")
    al = lambda r, a: r[k[a]] if k[a] is not None and k[a] < len(r) else None  # noqa: E731
    return [Soru(str(al(r, "id") or f"S{i}"), tarih_coz(al(r, "tarih")), str(al(r, "kullanici") or ""),
                 str(al(r, "departman") or ""), str(al(r, "metin") or "").strip())
            for i, r in enumerate(s[1:], 1) if str(al(r, "metin") or "").strip()]


def parola_maskele(metin: str) -> str:
    return PAROLA.sub(lambda m: (m.group(1) or m.group(3)) + "[GİZLİ]", metin)


def on_isle(sorular: list[Soru]) -> None:
    for s in sorular:
        if GUVENLIK.search(s.metin):
            s.etiketler.append("Güvenlik")
        if ARIZA.search(s.metin):
            s.etiketler.append("Arıza olabilir")
        if PAROLA.search(s.metin):
            s.etiketler.append("Soruda parola var")


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

SEMA = {
    "type": "object",
    "properties": {
        "cevaplar": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "id": {"type": "string"},
                    "kapsam": {"type": "string", "enum": ["tam", "kismi", "yok"]},
                    "cevap": {"type": "string"},
                    "kaynaklar": {"type": "array", "items": {"type": "string"}},
                    "takip_sorusu": {"type": "string"},
                    "insan_gerekli": {"type": "boolean"},
                    "gerekce": {"type": "string"},
                },
                "required": ["id", "kapsam", "cevap", "kaynaklar", "takip_sorusu", "insan_gerekli", "gerekce"],
                "additionalProperties": False,
            },
        },
    },
    "required": ["cevaplar"],
    "additionalProperties": False,
}


def soru_xml(s: Soru, bolumler: dict[str, Bolum]) -> str:
    kaynak = "\n".join(f'<kaynak id="{bid}" belge="{bolumler[bid].kaynak}" baslik="{bolumler[bid].baslik}">\n'
                       f"{llm.maskele(bolumler[bid].metin)}\n</kaynak>" for bid, _ in s.eslesme) or "<kaynak_yok/>"
    ek = f' departman="{s.departman}"' if s.departman else ""
    return (f'<soru id="{s.id}"{ek}>\n<metin>{llm.maskele(parola_maskele(s.metin[:SORU_SINIRI]))}</metin>\n'
            f"<ipuclari>{', '.join(s.etiketler) or '-'}</ipuclari>\n<kaynaklar>\n{kaynak}\n</kaynaklar>\n</soru>")


def cevapla(sorular: list[Soru], bolumler: list[Bolum]) -> dict[str, dict]:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    bmap = {b.id: b for b in bolumler}
    sonuc = {}
    for i in range(0, len(sorular), PAKET):
        parca = sorular[i:i + PAKET]
        mesaj = "<sorular>\n" + "\n".join(soru_xml(s, bmap) for s in parca) + "\n</sorular>"
        yanit = llm.json_iste(sistem, mesaj, SEMA)
        gecerli = {s.id for s in parca}
        for c in yanit.get("cevaplar", []):
            if c.get("id") in gecerli:
                sonuc[c["id"]] = c
    return sonuc


def denetle(s: Soru, c: dict, guven_esik: float) -> list[str]:
    """Kod kuralları modelin üstündedir. Uygulanan düzeltmeleri döndürür."""
    notlar = []
    verilen = {bid for bid, _ in s.eslesme}
    gecersiz = [k for k in c["kaynaklar"] if k not in verilen]
    if gecersiz:
        notlar.append(f"verilmeyen kaynak gösterildi ({', '.join(gecersiz)}), çıkarıldı")
        c["kaynaklar"] = [k for k in c["kaynaklar"] if k in verilen]
    metinde = set(re.findall(r"\[(K\d+)\]", c["cevap"]))
    if metinde - verilen:
        notlar.append("cevap metninde verilmeyen kaynak etiketi var")
        c["insan_gerekli"] = True
    if c["kapsam"] != "yok" and not c["kaynaklar"]:
        c["kapsam"] = "yok"
        notlar.append("kaynak gösterilmedi → kapsam yok sayıldı")
    if PAROLA_ISTEME.search(c["cevap"]):
        c["insan_gerekli"] = True
        notlar.append("taslak kullanıcıdan parola istiyor — gönderilmemeli")
    if "Güvenlik" in s.etiketler and not c["insan_gerekli"]:
        c["insan_gerekli"] = True
        notlar.append("güvenlik olayı belirtisi → Bilgi Güvenliği'ne yönlendirin (kural)")
    if "Arıza olabilir" in s.etiketler:
        notlar.append("arıza olabilir → destek talebi açılması önerilir")
    if c["kapsam"] == "yok":
        c["insan_gerekli"] = True
    en_iyi = max((p for _, p in s.eslesme), default=0.0)
    c["guven"] = ("yüksek" if c["kapsam"] == "tam" and en_iyi >= guven_esik and not c["insan_gerekli"] else
                  "orta" if c["kapsam"] in ("tam", "kismi") else "düşük")
    return notlar


def bosluklar(sorular: list[Soru], sonuc: dict[str, dict]) -> tuple[list[Soru], list[tuple[str, int]]]:
    """Bilgi bankasının cevaplayamadığı sorular ve bu sorularda en sık geçen kökler."""
    acik = [s for s in sorular if sonuc.get(s.id, {}).get("kapsam", "yok") != "tam"]
    sayac = Counter(t for s in acik for t in set(kokler(s.metin)))
    return acik, [(t, n) for t, n in sayac.most_common(15) if n >= 1]


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

KAPSAM_AD = {"tam": "Tam", "kismi": "Kısmi", "yok": "Yok"}
GUVEN_DOLGU = {"yüksek": PatternFill("solid", fgColor="E3F5E1"), "orta": PatternFill("solid", fgColor="FFF4CE"),
               "düşük": PatternFill("solid", fgColor="FDE2E1")}
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
ONAY_DOLGU = PatternFill("solid", fgColor="FFF4CE")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, satir=1):
    for h in ws[satir]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI


def rapor_yaz(cikti: Path, sorular: list[Soru], bolumler: list[Bolum], sonuc: dict, notlar: dict) -> None:
    bmap = {b.id: b for b in bolumler}
    wb = Workbook()
    ws = wb.active
    ws.title = "Cevaplar"
    ws.append(["Soru No", "Tarih", "Kullanıcı", "Departman", "Soru", "Kural Etiketleri", "Kapsam", "Güven", "Cevap Taslağı",
               "Dayandığı Bölümler", "Takip Sorusu", "İnsan Gerekli", "Gerekçe / Kod Notları", "Onay"])
    _baslik(ws)
    for j, w in enumerate((9, 15, 18, 14, 50, 18, 8, 8, 80, 40, 30, 9, 45, 12), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for s in sorular:
        c = sonuc.get(s.id)
        if c:
            dayanak = "\n".join(f"[{k}] {bmap[k].kaynak} › {bmap[k].baslik}" for k in c["kaynaklar"])
            gerekce = "; ".join(x for x in [c["gerekce"], *notlar.get(s.id, [])] if x)
            ws.append([s.id, s.tarih, s.kullanici, s.departman, s.metin, ", ".join(s.etiketler), KAPSAM_AD[c["kapsam"]],
                       c["guven"].capitalize(), c["cevap"], dayanak, c["takip_sorusu"], "EVET" if c["insan_gerekli"] else "Hayır",
                       gerekce, ""])
            ws.cell(ws.max_row, 8).fill = GUVEN_DOLGU[c["guven"]]
            for col in (7, 9, 10, 11):
                ws.cell(ws.max_row, col).fill = AI_DOLGU
            if c["insan_gerekli"]:
                ws.cell(ws.max_row, 12).fill = GUVEN_DOLGU["düşük"]
        else:
            ws.append([s.id, s.tarih, s.kullanici, s.departman, s.metin, ", ".join(s.etiketler), "", "Düşük",
                       "AI yanıt vermedi", "", "", "EVET", "Elle cevaplayın", ""])
        ws.cell(ws.max_row, 2).number_format = "DD.MM.YYYY HH:MM"
        ws.cell(ws.max_row, 14).fill = ONAY_DOLGU
        for h in ws[ws.max_row]:
            h.alignment = UST
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    e = wb.create_sheet("Kaynak Eşleşmeleri")
    e.append(["Soru No", "Sıra", "Bölüm", "Belge", "Başlık", "Arama Puanı", "Model Kullandı"])
    _baslik(e)
    for s in sorular:
        kullanilan = set(sonuc.get(s.id, {}).get("kaynaklar", []))
        for n, (bid, p) in enumerate(s.eslesme, 1):
            e.append([s.id, n, bid, bmap[bid].kaynak, bmap[bid].baslik, round(p, 2), "Evet" if bid in kullanilan else ""])
        if not s.eslesme:
            e.append([s.id, "-", "", "", "Eşleşen bölüm yok", 0, ""])
    for j, w in enumerate((9, 6, 7, 22, 50, 12, 14), 1):
        e.column_dimensions[get_column_letter(j)].width = w

    acik, kokler_ = bosluklar(sorular, sonuc)
    g = wb.create_sheet("Bilgi Bankası Boşlukları")
    g.append(["Bilgi bankasının tam cevaplayamadığı sorular — yeni madde yazılması veya mevcut maddenin genişletilmesi önerilir"])
    g["A1"].font = Font(bold=True)
    g.append(["Soru No", "Soru", "Kapsam", "Model Gerekçesi"])
    _baslik(g, 2)
    for s in acik:
        c = sonuc.get(s.id, {})
        g.append([s.id, s.metin, KAPSAM_AD.get(c.get("kapsam", "yok"), "Yok"), c.get("gerekce", "")])
        for h in g[g.max_row]:
            h.alignment = UST
    g.append([])
    g.append(["Bu sorularda sık geçen kelime kökleri", "Soru sayısı"])
    _baslik(g, g.max_row)
    for t, n in kokler_:
        g.append([t, n])
    g.column_dimensions["A"].width = 30
    g.column_dimensions["B"].width = 80
    g.column_dimensions["C"].width = 9
    g.column_dimensions["D"].width = 50

    o = wb.create_sheet("Özet", 0)
    o.append(["Sık sorulan sorular — cevap taslakları"])
    o["A1"].font = Font(bold=True, size=13)
    kapsam = Counter(KAPSAM_AD[c["kapsam"]] for c in sonuc.values())
    guven = Counter(c["guven"].capitalize() for c in sonuc.values())
    for s in [("Toplam soru", len(sorular)), ("Bilgi bankası", f"{len({b.kaynak for b in bolumler})} belge, {len(bolumler)} bölüm"),
              ("Kapsam: tam / kısmi / yok", f"{kapsam['Tam']} / {kapsam['Kısmi']} / {kapsam['Yok'] + len(sorular) - len(sonuc)}"),
              ("Güven: yüksek / orta / düşük", f"{guven['Yüksek']} / {guven['Orta']} / {guven['Düşük'] + len(sorular) - len(sonuc)}"),
              ("İnsan gerekli", sum(1 for c in sonuc.values() if c["insan_gerekli"]) + len(sorular) - len(sonuc)),
              ("Model", llm.kullanim_ozeti()), ("Oluşturulma", datetime.now().strftime("%d.%m.%Y %H:%M")),
              ("Önemli", "Cevaplar taslaktır ve otomatik gönderilmez. Model yalnız aramayla bulunan bilgi bankası bölümlerini "
                         "görür; kaynak göstermeyen cevaplar 'kapsam yok' sayılır. Güvenlik olayı belirtisi olan sorular ve "
                         "parola isteyen taslaklar insana yönlendirilir (kod kuralı).")]:
        o.append(list(s))
    o.column_dimensions["A"].width = 30
    o.column_dimensions["B"].width = 100
    for satir in o.iter_rows():
        for h in satir:
            h.alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(sorular: list[Soru], bilgi: Path, cikti: Path, k: int = 4, esik: float = 1.0, guven_esik: float = 4.0,
             evet: bool = False) -> dict:
    if not sorular:
        raise llm.LLMHatasi("Cevaplanacak soru yok.")
    bolumler = bilgi_bankasi_oku(bilgi)
    on_isle(sorular)
    eslestir(sorular, bolumler, k, esik)
    gonderilen = {bid for s in sorular for bid, _ in s.eslesme}
    etiket = Counter(e for s in sorular for e in s.etiketler)
    print(f"[OK] {len(sorular)} soru · bilgi bankası {len(bolumler)} bölüm · eşleşmesi olmayan "
          f"{sum(1 for s in sorular if not s.eslesme)} soru" + ("".join(f" · {e}: {n}" for e, n in etiket.items())))
    llm.onay_al(f"{len(sorular)} soru (e-posta/telefon/IBAN/TCKN/parola maskeli) ve bilgi bankasının bu sorularla eşleşen "
                f"{len(gonderilen)}/{len(bolumler)} bölümü gönderilecek (eşleşmeyen bölümler gönderilmez).", evet)
    sonuc = cevapla(sorular, bolumler)
    notlar = {s.id: denetle(s, sonuc[s.id], guven_esik) for s in sorular if s.id in sonuc}
    rapor_yaz(cikti, sorular, bolumler, sonuc, notlar)
    return {"sorular": sorular, "bolumler": bolumler, "sonuc": sonuc, "notlar": notlar}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    ornek = BURASI / "ornek_veri"
    p = argparse.ArgumentParser(description="Kullanıcı sorularına şirket bilgi bankasından kaynaklı cevap taslağı üretir.")
    p.add_argument("--girdi", type=Path, help="Sorular (.xlsx/.csv: Soru No, Tarih, Kullanıcı, Departman, Soru)")
    p.add_argument("--soru", help="Tek bir soru (dosya yerine)")
    p.add_argument("--bilgi", type=Path, default=ornek / "bilgi_bankasi", help="Bilgi bankası dosyası veya klasörü (.md/.txt/.docx/.pdf)")
    p.add_argument("--k", type=int, default=4, help="Soru başına modele gönderilecek en fazla bölüm (varsayılan 4)")
    p.add_argument("--esik", type=float, default=1.0, help="Bölümün gönderilmesi için en düşük arama puanı (varsayılan 1)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "sss_cevaplari.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        if a.soru:
            sorular = [Soru("S1", datetime.now(), "", "", a.soru.strip())]
        else:
            sorular = sorulari_oku(a.girdi or ornek / "sorular.csv")
        s = calistir(sorular, a.bilgi, a.cikti, a.k, a.esik, evet=a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    if a.soru and "S1" in s["sonuc"]:
        c = s["sonuc"]["S1"]
        print(f"\n{c['cevap']}\n\n[i] Kapsam: {KAPSAM_AD[c['kapsam']]} · güven: {c['guven']} · kaynak: {', '.join(c['kaynaklar']) or '-'}")
    insan = sum(1 for x in s["sonuc"].values() if x["insan_gerekli"])
    print(f"[OK] {len(s['sonuc'])}/{len(s['sorular'])} soru cevaplandı · insan gerekli: {insan}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
