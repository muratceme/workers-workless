"""
Ürün Sorusu Cevap Taslağı — Workers / Workless AI Agent
E-Ticaret › Müşteri Deneyimi › Müşteri Hizmetleri Temsilcisi

1. Pazaryeri ürün soruları (Excel/CSV) okunur; her soru ürün koduyla ürün kartına (özellikler tablosu) bağlanır.
   Şirket politikalarından (kargo, iade, garanti…) soruyla ilgili bölümler BM25 aramasıyla bulunur.
2. Model YALNIZ ürün kartı ve verilen politika bölümlerine dayanarak kısa, kibar bir cevap taslağı yazar;
   bilgi yoksa uydurmaz, "insan gerekli" işaretler.
3. Kod denetler (modelin üstünde):
   - Cevapta telefon, e-posta, web adresi, sosyal medya veya platform dışı yönlendirme olamaz (pazaryeri kuralları).
   - Cevaptaki ölçü/ağırlık/sıcaklık/yüzde gibi sayılar ürün kartında veya politikada geçmiyorsa
     "doğrulanamayan bilgi" olarak işaretlenir.
   - Sağlık/tedavi iddiası ve kesin teslim tarihi sözü işaretlenir; ürün kartı bulunamayan sorular insana gider.
4. Sonuç Excel'e yazılır; her taslak için 'Onay' sütunu vardır. Hiçbir cevap otomatik gönderilmez.

Kullanım:
    python agent.py                                                # örnek sorularla dener
    python agent.py --sorular sorular.xlsx --urunler urunler.xlsx --politika politikalar.md
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm
import sss_cekirdek as sss

BURASI = Path(__file__).resolve().parent
PAKET = 15

ILETISIM = re.compile(r"(\b0?\s?5\d{2}[\s.-]?\d{3}[\s.-]?\d{2}[\s.-]?\d{2}\b|\b0?\s?850[\s.-]?\d{3}[\s.-]?\d{2}[\s.-]?\d{2}\b|"
                      r"[\w.+-]+@[\w-]+\.[\w.]+|https?://|www\.|\.com\b|\.com\.tr\b|whatsapp|instagram|telegram|"
                      r"mağazamıza gel|sitemizden|web sitemiz|dm\b|direkt mesaj)", re.I)
SAGLIK = re.compile(r"(tedavi eder|iyileştir|hastalığ|şifa|kanser|zayıflat|kilo verdir|bağışıklığı güçlendir|"
                    r"yan etkisi yok|%100 doğal ve zararsız)", re.I)
TESLIM_SOZU = re.compile(r"(yarın (?:elinizde|teslim|kargoda)|bugün kargoda|\d+ (?:saat|gün) içinde (?:elinizde|teslim))", re.I)
SAYI = re.compile(r"(\d+(?:[.,]\d+)?)\s*(cm|mm|m\b|metre|kg|gr|g\b|g/m²|°c|derece|lt|litre|ml|w\b|watt|%|yüzde|gün|ay|yıl|adet)", re.I)
SORU_ILETISIM = re.compile(r"(numara(nızı|nız)|telefon|whatsapp|instagram|iletişim bilgi|mail adres)", re.I)


def sayi_kumesi(metin: str) -> set[str]:
    """Metindeki tüm sayılar (birimden bağımsız, '1,5' ve '1.5' aynı)."""
    return {n.replace(",", ".").rstrip("0").rstrip(".") if "." in n.replace(",", ".") else n
            for n in re.findall(r"\d+(?:[.,]\d+)?", metin)}


# ----------------------------------------------------------------------------
# Okuma
# ----------------------------------------------------------------------------

def _tablo(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = belge.txt_oku(yol)
        ilk = "\n".join(metin.splitlines()[:5])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    satirlar = [r for r in satirlar if any(c not in (None, "") for c in r)]
    if not satirlar:
        raise llm.LLMHatasi(f"{yol.name}: dosya boş.")
    return satirlar


def urunleri_oku(yol: Path) -> dict[str, dict[str, str]]:
    satirlar = _tablo(yol)
    bas = [str(x or "").strip() for x in satirlar[0]]
    kb = [sss.katla(x) for x in bas]
    i_kod = next((kb.index(sss.katla(a)) for a in ("stok kodu", "ürün kodu", "model kodu", "barkod", "kod") if sss.katla(a) in kb), 0)
    urunler = {}
    for r in satirlar[1:]:
        if r[i_kod] in (None, ""):
            continue
        urunler[str(r[i_kod]).strip()] = {bas[j]: str(v).strip() for j, v in enumerate(r) if j < len(bas) and v not in (None, "")}
    return urunler


def sorulari_oku(yol: Path) -> list[dict]:
    satirlar = _tablo(yol)
    kb = [sss.katla(x) for x in satirlar[0]]
    bul = lambda *a: next((kb.index(sss.katla(x)) for x in a if sss.katla(x) in kb), None)  # noqa: E731
    i = {"no": bul("soru no", "no", "id"), "tarih": bul("tarih", "soru tarihi"), "kanal": bul("kanal", "pazaryeri", "platform"),
         "urun": bul("ürün kodu", "stok kodu", "model kodu", "barkod"), "soru": bul("soru", "soru metni", "mesaj")}
    if i["soru"] is None:
        raise llm.LLMHatasi(f"Soru dosyasında 'Soru' sütunu bulunamadı. Başlıklar: {satirlar[0]}")
    al = lambda r, k: r[i[k]] if i[k] is not None and i[k] < len(r) else None  # noqa: E731
    return [{"id": str(al(r, "no") or f"S{n}"), "tarih": sss.tarih_coz(al(r, "tarih")), "kanal": str(al(r, "kanal") or ""),
             "urun": str(al(r, "urun") or "").strip(), "soru": str(al(r, "soru") or "").strip(), "etiketler": [], "eslesme": []}
            for n, r in enumerate(satirlar[1:], 1) if str(al(r, "soru") or "").strip()]


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

SEMA = {
    "type": "object",
    "properties": {
        "cevaplar": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "id": {"type": "string"},
                    "kapsam": {"type": "string", "enum": ["tam", "kismi", "yok"]},
                    "cevap": {"type": "string"},
                    "dayanak": {"type": "array", "items": {"type": "string"}},
                    "insan_gerekli": {"type": "boolean"},
                    "gerekce": {"type": "string"},
                },
                "required": ["id", "kapsam", "cevap", "dayanak", "insan_gerekli", "gerekce"],
                "additionalProperties": False,
            },
        },
    },
    "required": ["cevaplar"],
    "additionalProperties": False,
}


def soru_xml(s: dict, urun: dict | None, bolumler: dict) -> str:
    kart = "\n".join(f"<ozellik ad=\"{k}\">{v}</ozellik>" for k, v in (urun or {}).items()) or "<urun_bulunamadi/>"
    pol = "\n".join(f'<kaynak id="{bid}" baslik="{bolumler[bid].baslik}">\n{bolumler[bid].metin}\n</kaynak>'
                    for bid, _ in s["eslesme"]) or "<kaynak_yok/>"
    return (f'<soru id="{s["id"]}" kanal="{s["kanal"]}">\n<metin>{llm.maskele(s["soru"][:1500])}</metin>\n'
            f"<ipuclari>{', '.join(s['etiketler']) or '-'}</ipuclari>\n<urun_karti>\n{kart}\n</urun_karti>\n"
            f"<politikalar>\n{pol}\n</politikalar>\n</soru>")


def denetle(s: dict, c: dict, urun: dict | None, bolumler: dict, azami: int) -> list[str]:
    notlar = []
    gecerli = {bid for bid, _ in s["eslesme"]} | set((urun or {}).keys())
    c["dayanak"] = [d for d in c["dayanak"] if d in gecerli]
    if ILETISIM.search(c["cevap"]):
        c["insan_gerekli"] = True
        notlar.append("cevapta iletişim bilgisi/platform dışı yönlendirme var — pazaryeri kurallarına aykırı, düzeltin")
    if SAGLIK.search(c["cevap"]):
        c["insan_gerekli"] = True
        notlar.append("sağlık/tedavi iddiası içeriyor — kaldırın")
    if TESLIM_SOZU.search(c["cevap"]):
        c["insan_gerekli"] = True
        notlar.append("kesin teslim sözü içeriyor — politika ve stok durumuna göre kontrol edin")
    kaynak_metin = " ".join(f"{k} {v}" for k, v in (urun or {}).items()) + " " + \
        " ".join(bolumler[b].metin for b, _ in s["eslesme"]) + " " + s["soru"]
    bilinen = sayi_kumesi(kaynak_metin)
    supheli = sorted({m.group(0).strip() for m in SAYI.finditer(c["cevap"])
                      if (m.group(1).replace(",", ".").rstrip("0").rstrip(".") if "." in m.group(1).replace(",", ".") else m.group(1))
                      not in bilinen})
    if supheli:
        c["insan_gerekli"] = True
        notlar.append("ürün kartı/politikada olmayan sayısal bilgi: " + ", ".join(supheli))
    if urun is None:
        c["insan_gerekli"] = True
        notlar.append("ürün kartı bulunamadı")
    if len(c["cevap"]) > azami:
        notlar.append(f"cevap {len(c['cevap'])} karakter (sınır {azami})")
    if c["kapsam"] == "yok":
        c["insan_gerekli"] = True
    return notlar


def calistir(sorular_yolu: Path, urunler_yolu: Path, politika_yolu: Path, cikti: Path, azami: int = 500,
             evet: bool = False) -> dict:
    sorular = sorulari_oku(sorular_yolu)
    if not sorular:
        raise llm.LLMHatasi("Cevaplanacak soru yok.")
    urunler = urunleri_oku(urunler_yolu)
    bolumler = sss.bilgi_bankasi_oku(politika_yolu)
    bmap = {b.id: b for b in bolumler}
    motor = sss.BM25([b.kok for b in bolumler])
    for s in sorular:
        if s["urun"] not in urunler:
            s["etiketler"].append("Ürün kartı yok")
        if SORU_ILETISIM.search(s["soru"]):
            s["etiketler"].append("İletişim bilgisi istiyor")
        if SAGLIK.search(s["soru"]):
            s["etiketler"].append("Sağlık sorusu")
        puan = motor.puan(sss.kokler(s["soru"]))
        sira = sorted(range(len(bolumler)), key=lambda i: -puan[i])
        s["eslesme"] = [(bolumler[i].id, puan[i]) for i in sira[:2] if puan[i] >= 1.0]
    etiket = Counter(e for s in sorular for e in s["etiketler"])
    print(f"[OK] {len(sorular)} soru · {len(urunler)} ürün kartı · {len(bolumler)} politika bölümü"
          + "".join(f" · {e}: {n}" for e, n in etiket.items()))
    llm.onay_al(f"{len(sorular)} soru (kişisel veriler maskeli), ilgili ürün kartları ve politika bölümleri gönderilecek.", evet)
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    sonuc = {}
    for i in range(0, len(sorular), PAKET):
        parca = sorular[i:i + PAKET]
        mesaj = "<sorular>\n" + "\n".join(soru_xml(s, urunler.get(s["urun"]), bmap) for s in parca) + "\n</sorular>"
        yanit = llm.json_iste(sistem, mesaj, SEMA)
        gecerli = {s["id"] for s in parca}
        for c in yanit.get("cevaplar", []):
            if c.get("id") in gecerli:
                sonuc[c["id"]] = c
    notlar = {s["id"]: denetle(s, sonuc[s["id"]], urunler.get(s["urun"]), bmap, azami) for s in sorular if s["id"] in sonuc}
    _rapor(sorular, sonuc, notlar, urunler, bmap, cikti)
    return {"sorular": sorular, "sonuc": sonuc, "notlar": notlar}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
ONAY_DOLGU = PatternFill("solid", fgColor="FFF4CE")
INSAN_DOLGU = PatternFill("solid", fgColor="FDE2E1")
UST = Alignment(vertical="top", wrap_text=True)
KAPSAM_AD = {"tam": "Tam", "kismi": "Kısmi", "yok": "Yok"}


def _rapor(sorular, sonuc, notlar, urunler, bmap, cikti):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cevaplar"
    ws.append(["Soru No", "Tarih", "Kanal", "Ürün Kodu", "Ürün", "Soru", "Kural Etiketleri", "Kapsam", "Cevap Taslağı",
               "Dayandığı Bilgi", "İnsan Gerekli", "Gerekçe / Kod Notları", "Onay"])
    for h in ws[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for s in sorular:
        c = sonuc.get(s["id"])
        u = urunler.get(s["urun"], {})
        ad = " ".join(v for k, v in u.items() if sss.katla(k) in ("marka", "urun tipi", "urun adi", "ad"))[:60]
        if c:
            dayanak = ", ".join(f"[{d}] {bmap[d].baslik}" if d in bmap else d for d in c["dayanak"])
            ws.append([s["id"], s["tarih"], s["kanal"], s["urun"], ad, s["soru"], ", ".join(s["etiketler"]), KAPSAM_AD[c["kapsam"]],
                       c["cevap"], dayanak, "EVET" if c["insan_gerekli"] else "Hayır",
                       "; ".join(x for x in [c["gerekce"], *notlar.get(s["id"], [])] if x), ""])
            for col in (8, 9, 10):
                ws.cell(ws.max_row, col).fill = AI_DOLGU
            if c["insan_gerekli"]:
                ws.cell(ws.max_row, 11).fill = INSAN_DOLGU
        else:
            ws.append([s["id"], s["tarih"], s["kanal"], s["urun"], ad, s["soru"], ", ".join(s["etiketler"]), "", "AI yanıt vermedi",
                       "", "EVET", "Elle cevaplayın", ""])
        ws.cell(ws.max_row, 2).number_format = "DD.MM.YYYY HH:MM"
        ws.cell(ws.max_row, 13).fill = ONAY_DOLGU
        for h in ws[ws.max_row]:
            h.alignment = UST
    for j, w in enumerate((9, 15, 12, 11, 24, 45, 18, 8, 70, 28, 9, 50, 12), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    o = wb.create_sheet("Özet", 0)
    o.append(["Ürün soruları — cevap taslakları"])
    o["A1"].font = Font(bold=True, size=13)
    kapsam = Counter(KAPSAM_AD[c["kapsam"]] for c in sonuc.values())
    for satir in [("Toplam soru", len(sorular)),
                  ("Kapsam: tam / kısmi / yok", f"{kapsam['Tam']} / {kapsam['Kısmi']} / {kapsam['Yok'] + len(sorular) - len(sonuc)}"),
                  ("İnsan gerekli", sum(1 for c in sonuc.values() if c["insan_gerekli"]) + len(sorular) - len(sonuc)),
                  ("En çok soru alan ürünler", ", ".join(f"{k} ({n})" for k, n in Counter(s["urun"] for s in sorular).most_common(5))),
                  ("Model", llm.kullanim_ozeti()), ("Oluşturulma", datetime.now().strftime("%d.%m.%Y %H:%M")),
                  ("Önemli", "Cevaplar taslaktır ve otomatik gönderilmez. Model yalnız ürün kartını ve eşleşen politika bölümlerini "
                             "görür. Kod; iletişim bilgisi/platform dışı yönlendirme, sağlık iddiası, kesin teslim sözü ve ürün kartında "
                             "olmayan sayısal bilgileri yakalar. Sık sorulan bilgileri ürün kartına/açıklamasına eklemeyi düşünün.")]:
        o.append(list(satir))
    o.column_dimensions["A"].width = 28
    o.column_dimensions["B"].width = 100
    for r in o.iter_rows():
        for h in r:
            h.alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    ornek = BURASI / "ornek_veri"
    p = argparse.ArgumentParser(description="Pazaryeri ürün sorularına ürün kartı ve şirket politikalarından cevap taslağı hazırlar.")
    p.add_argument("--sorular", type=Path, default=ornek / "sorular.csv", help="Sorular (.xlsx/.csv): Soru No, Tarih, Kanal, Ürün Kodu, Soru")
    p.add_argument("--urunler", type=Path, default=ornek / "urunler.csv", help="Ürün kartları (.xlsx/.csv): ilk sütun ürün kodu, diğerleri özellikler")
    p.add_argument("--politika", type=Path, default=ornek / "politikalar.md", help="Kargo, iade, garanti politikaları (.md/.txt/.docx/.pdf ya da klasör)")
    p.add_argument("--azami", type=int, default=500, help="Cevap için karakter sınırı uyarısı (varsayılan 500)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "urun_sorulari.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.sorular, a.urunler, a.politika, a.cikti, a.azami, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    insan = sum(1 for x in s["sonuc"].values() if x["insan_gerekli"])
    print(f"[OK] {len(s['sonuc'])}/{len(s['sorular'])} soru cevaplandı · insan gerekli: {insan}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
