"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

V = KLASOR / "ornek_veri"
PLAN, FOTO = V / "planogram.csv", V / "raf_foto.jpg"
CAGRILAR = []


def g(raf, sira, urun, yuz, etiket, fiyat, gorulen="", guven="Yüksek"):
    return {"raf": raf, "sira": sira, "urun": urun, "gorulen": gorulen or "ürün", "yuz_sayisi": yuz, "fiyat_etiketi": etiket, "etiket_fiyati": fiyat, "guven": guven}


def gozlem(dogru=True) -> dict:
    t = {"gozlemler": [g(1, 1, "8690000000011", 3, "Var", "189,90"), g(1, 2, "8690000000035", 2, "Var", "64,50"), g(1, 3, "8690000000028", 2, "Yok", ""),
                       g(2, 1, "8690000000042", 4, "Var", "79,90"), g(2, 2, "8690000000059", 1, "Var", "199,00"),
                       g(3, 1, "8690000000073", 3, "Var", "69,90"), g(3, 2, "8690000000080", 3, "Var", "54,90", guven="Düşük"),
                       g(3, 3, agent.BILINMIYOR, 1, "Yok", "", "Sarı ambalajlı başka marka şeker")],
         "bos_alanlar": [{"raf": 2, "sira": 3}], "genel_not": "Fotoğraf net; tüm raflar görünüyor."}
    if not dogru:
        t["gozlemler"] = t["gozlemler"][:5] + [g(1, 1, "8690000000099", 2, "Var", ""), g(4, 1, "8690000000073", 3, "Var", "69,90")]
    return t


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000, gorseller=None):
    CAGRILAR.append((kullanici, gorseller))
    return gozlem("<duzeltme>" in kullanici)


class KodTesti(unittest.TestCase):
    def test_okuma(self):
        p = agent.planogram_oku(PLAN)
        self.assertEqual((len(p), p[4]["fiyat"], p[0]["yuz"]), (8, 214.0, 3))
        f = agent.foto_hazirla([FOTO])
        self.assertEqual((f[0][0], f[0][1][:2]), ("image/jpeg", b"\xff\xd8"))

    def test_denetim_ve_karsilastirma(self):
        p = agent.planogram_oku(PLAN)
        self.assertEqual(agent.denetle(gozlem(), p), [])
        s = agent.denetle(gozlem(False), p)
        for parca in ("'8690000000099' planogramda yok", "aynı konuma iki ürün", "fiyat yazılmamış", "Raf 4", "Raf 3: gözlem yok"):
            self.assertTrue(any(parca in x for x in s), (parca, s))
        self.assertEqual(len(s), 5, s)
        satir, fazla = agent.karsilastir(gozlem(), p)
        d = {x["ad"]: (x["durum"], x["sorunlar"]) for x in satir}
        self.assertEqual(d["Örnek Çay 500 g"], ("Yerinde", []))
        self.assertEqual(d["Örnek Çay 1 kg"], ("Yanlış yer", ["raf 1 / sıra 3 konumunda", "fiyat etiketi yok"]))
        self.assertEqual(d["Deneme Bitki Çayı 20'li"][0], "Yanlış yer")
        self.assertEqual(d["Deneme Filtre Kahve 250 g"], ("Yerinde", ["yüz sayısı 1 / 3", "etiket fiyatı 199,00, planogram 214,00"]))
        self.assertEqual(d["Deneme Hazır Kahve 200 g"], ("Eksik", []))
        self.assertEqual((len(fazla), [x["elle"] for x in satir].count(True)), (1, 1))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        CAGRILAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "u.xlsx"
            s = agent.calistir(PLAN, [FOTO], cikti, evet=True)
            self.assertEqual((len(CAGRILAR), s["kalan"], s["tur"], s["uyumlu"]), (2, [], 1, 4))
            self.assertTrue(all(len(c[1]) == 1 and c[1][0][0] == "image/jpeg" for c in CAGRILAR))
            self.assertIn("raf 2 | sıra 3 | 8690000000066 | Deneme Hazır Kahve 200 g | 2 yüz", CAGRILAR[0][0])
            ws = load_workbook(cikti)["Uyum"]
            self.assertEqual((ws["D7"].value, ws["G7"].value), ("Deneme Hazır Kahve 200 g", "Eksik"))
            self.assertEqual(load_workbook(cikti)["Planogram Dışı"]["C2"].value, "Sarı ambalajlı başka marka şeker")

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(PLAN, [FOTO], Path(tmp) / "u.xlsx", tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 5)
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--evet", "--foto", str(Path(tmp) / "yok.jpg")]), 1)


if __name__ == "__main__":
    unittest.main()
