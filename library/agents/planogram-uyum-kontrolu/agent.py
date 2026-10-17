"""
Planogram Uyum Kontrolü — Gerisi Bende AI Agent
Görsel Düzenleme › Görsel Düzenleme Uzmanı

1. Kod planogram listesini (raf, sıra, barkod, ürün adı, yüz sayısı, fiyat) ve raf fotoğraf(lar)ını okur. Fotoğraflar
   uzun kenarı en fazla 1568 piksel olacak şekilde küçültülür.
2. Model (görsel destekli) fotoğrafta **gördüğünü** yazar: her raf ve sıradaki ürün (planogramdaki barkod ya da
   "BILINMIYOR"), görünen yüz sayısı, fiyat etiketi (Var / Yok / Okunamıyor) ve etiketteki fiyat, boş alanlar. Uyum
   kararını model vermez.
3. Kod modelin gözlemlerini denetler (en fazla --tur kez geri gönderilir):
   - Ürün planogramda var mı ya da "BILINMIYOR" mu; raf numarası planogram aralığında mı; aynı konuma iki ürün yazılmış mı
   - Etiket "Var" ise fiyat okunabilir mi; yüz sayısı en az 1 mi; her raf için gözlem ya da boş alan var mı
4. Kod planogramla karşılaştırır: Yerinde / Yanlış yer / Eksik, yüz sayısı eksiği, fiyat etiketi yok, etiket fiyatı
   planogramdan farklı, planogramda olmayan ürün. Düşük güvenli gözlemler "elle kontrol" diye işaretlenir.

Kullanım:
    python agent.py                                                          # örnek (kurgusal) sıcak içecek reyonu
    python agent.py --planogram reyon12.xlsx --foto raf1.jpg raf2.jpg
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, UnidentifiedImageError

import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
PLAN = {"raf": ("Raf", "Raf No"), "sira": ("Sıra", "Pozisyon", "Sıra No"), "barkod": ("Barkod", "Ürün Kodu", "EAN"), "ad": ("Ürün Adı", "Ürün"),
        "yuz": ("Yüz Sayısı", "Facing", "Yüz"), "fiyat": ("Fiyat", "Satış Fiyatı", "Raf Fiyatı")}
BILINMIYOR = "BILINMIYOR"
ETIKET = ["Var", "Yok", "Okunamıyor"]
GUVEN = ["Yüksek", "Orta", "Düşük"]
I = {"type": "integer"}

SEMA = {
    "type": "object",
    "properties": {
        "gozlemler": {"type": "array", "items": {"type": "object", "properties": {
            "raf": I, "sira": I, "urun": S, "gorulen": S, "yuz_sayisi": I, "fiyat_etiketi": {"type": "string", "enum": ETIKET}, "etiket_fiyati": S,
            "guven": {"type": "string", "enum": GUVEN}},
            "required": ["raf", "sira", "urun", "gorulen", "yuz_sayisi", "fiyat_etiketi", "etiket_fiyati", "guven"], "additionalProperties": False}},
        "bos_alanlar": {"type": "array", "items": {"type": "object", "properties": {"raf": I, "sira": I}, "required": ["raf", "sira"], "additionalProperties": False}},
        "genel_not": S,
    },
    "required": ["gozlemler", "bos_alanlar", "genel_not"],
    "additionalProperties": False,
}


def fiyat_coz(x) -> float | None:
    if isinstance(x, (int, float)):
        return float(x)
    m = re.search(r"\d+(?:[.,]\d+)*", str(x or ""))
    if not m:
        return None
    try:
        return float(sayi_norm(m.group(0)))
    except ValueError:
        return None


def planogram_oku(yol: Path) -> list[dict]:
    out = []
    for r in kayitlar(yol, PLAN, ("raf", "sira", "barkod")):
        if not str(r["barkod"]).strip():
            continue
        try:
            raf, sira, yuz = int(float(r["raf"])), int(float(r["sira"])), int(float(r["yuz"] or 1))
        except ValueError:
            raise llm.LLMHatasi(f"{yol.name}: raf / sıra / yüz sayısı tam sayı olmalı ({r['barkod']}).")
        out.append({"raf": raf, "sira": sira, "barkod": str(r["barkod"]).strip(), "ad": str(r["ad"]), "yuz": yuz, "fiyat": fiyat_coz(r["fiyat"])})
    if not out:
        raise llm.LLMHatasi(f"{yol.name}: planogram boş.")
    return out


def foto_hazirla(yollar: list[Path], uzun: int = 1568) -> list[tuple[str, bytes]]:
    out = []
    for y in yollar:
        if not y.exists():
            raise llm.LLMHatasi(f"{y}: fotoğraf bulunamadı.")
        try:
            with Image.open(y) as im:
                im = im.convert("RGB")
                im.thumbnail((uzun, uzun))
                b = io.BytesIO()
                im.save(b, "JPEG", quality=85)
        except (UnidentifiedImageError, OSError) as h:
            raise llm.LLMHatasi(f"{y.name}: fotoğraf açılamadı ({h}).")
        out.append(("image/jpeg", b.getvalue()))
    return out


def denetle(t: dict, plan: list[dict]) -> list[str]:
    sorun = []
    barkodlar = {p["barkod"] for p in plan}
    raflar = {p["raf"] for p in plan}
    konum = {}
    for g in t["gozlemler"]:
        ad = f"Raf {g['raf']} / sıra {g['sira']}"
        if g["urun"] != BILINMIYOR and g["urun"] not in barkodlar:
            sorun.append(f"{ad}: '{g['urun']}' planogramda yok; planogramdaki barkodu ya da {BILINMIYOR} yazın")
        if g["urun"] == BILINMIYOR and len(g["gorulen"].strip()) < 3:
            sorun.append(f"{ad}: bilinmeyen ürün için görüneni tarif edin")
        if g["raf"] not in raflar:
            sorun.append(f"{ad}: planogramda {min(raflar)}-{max(raflar)} arası raf var")
        if (g["raf"], g["sira"]) in konum:
            sorun.append(f"{ad}: aynı konuma iki ürün yazılmış ({konum[(g['raf'], g['sira'])]} ve {g['urun']})")
        konum[(g["raf"], g["sira"])] = g["urun"]
        if g["yuz_sayisi"] < 1:
            sorun.append(f"{ad}: yüz sayısı en az 1 olmalı")
        if g["fiyat_etiketi"] == "Var" and fiyat_coz(g["etiket_fiyati"]) is None:
            sorun.append(f"{ad}: etiket 'Var' ama fiyat yazılmamış; okunamıyorsa 'Okunamıyor' seçin")
    for b in t["bos_alanlar"]:
        if (b["raf"], b["sira"]) in konum:
            sorun.append(f"Raf {b['raf']} / sıra {b['sira']}: hem ürün hem boş alan yazılmış")
    for r in sorted(raflar):
        if not any(g["raf"] == r for g in t["gozlemler"]) and not any(b["raf"] == r for b in t["bos_alanlar"]):
            sorun.append(f"Raf {r}: gözlem yok (raf fotoğrafta görünmüyorsa genel notta belirtin ve boş alan yazmayın)")
    return sorun


def karsilastir(t: dict, plan: list[dict], tolerans: float = 0.005) -> tuple[list[dict], list[dict]]:
    gor = {}
    for g in t["gozlemler"]:
        gor.setdefault(g["urun"], []).append(g)
    satirlar = []
    for p in sorted(plan, key=lambda x: (x["raf"], x["sira"])):
        gs = gor.get(p["barkod"], [])
        sorunlar = []
        if not gs:
            durum, g = "Eksik", None
        else:
            g = next((x for x in gs if (x["raf"], x["sira"]) == (p["raf"], p["sira"])), None)
            durum = "Yerinde" if g else "Yanlış yer"
            g = g or gs[0]
            yuz = sum(x["yuz_sayisi"] for x in gs)
            if yuz < p["yuz"]:
                sorunlar.append(f"yüz sayısı {yuz} / {p['yuz']}")
            if g["fiyat_etiketi"] == "Yok":
                sorunlar.append("fiyat etiketi yok")
            elif g["fiyat_etiketi"] == "Okunamıyor":
                sorunlar.append("fiyat etiketi okunamadı")
            elif p["fiyat"] is not None and fiyat_coz(g["etiket_fiyati"]) is not None and abs(fiyat_coz(g["etiket_fiyati"]) - p["fiyat"]) > tolerans:
                sorunlar.append(f"etiket fiyatı {g['etiket_fiyati']}, planogram {p['fiyat']:.2f}".replace(".", ","))
            if durum == "Yanlış yer":
                sorunlar.insert(0, f"raf {g['raf']} / sıra {g['sira']} konumunda")
        satirlar.append(p | {"durum": durum, "sorunlar": sorunlar, "gozlem": g, "elle": bool(g and g["guven"] == "Düşük")})
    fazla = gor.get(BILINMIYOR, [])
    return satirlar, fazla


def uret(plan, fotolar, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    satir = "\n".join(f"raf {p['raf']} | sıra {p['sira']} | {p['barkod']} | {p['ad']} | {p['yuz']} yüz" for p in sorted(plan, key=lambda x: (x["raf"], x["sira"])))
    mesaj = f"<planogram>\n{satir}\n</planogram>\n<fotograf_sayisi>{len(fotolar)}</fotograf_sayisi>" + duzeltme_blogu(duzeltme, onceki)
    return llm.json_iste(sistem, mesaj, SEMA, gorseller=fotolar)


def rapor_yaz(cikti: Path, t, satirlar, fazla, kalan, tur) -> dict:
    wb = Workbook()
    ws = wb.active
    ws.title = "Uyum"
    baslik(ws, ["Raf", "Sıra", "Barkod", "Ürün", "Planogram Yüz", "Planogram Fiyat", "Durum", "Sorunlar", "Fotoğrafta Görülen", "Güven", "Elle Kontrol"],
           (6, 6, 16, 30, 10, 10, 12, 45, 30, 9, 10))
    for s in satirlar:
        g = s["gozlem"] or {}
        ws.append([s["raf"], s["sira"], s["barkod"], s["ad"], s["yuz"], s["fiyat"], s["durum"], "; ".join(s["sorunlar"]), g.get("gorulen", ""), g.get("guven", ""),
                   "Evet" if s["elle"] else ""])
        boya(ws, (7,), YESIL if s["durum"] == "Yerinde" and not s["sorunlar"] else (KIRMIZI if s["durum"] != "Yerinde" else SARI))
        boya(ws, (9,), AI_DOLGU)
    uyumlu = sum(1 for s in satirlar if s["durum"] == "Yerinde" and not s["sorunlar"])
    oran = uyumlu / len(satirlar) * 100
    son_not(ws, f"Uyum: {uyumlu}/{len(satirlar)} ürün tam uyumlu (%{oran:.0f}). Model notu: {t['genel_not']}\nÖnemli: Fotoğraf yorumu modelden gelir; "
                "'Elle Kontrol' işaretli ve kırmızı satırları rafta doğrulayın.", kalan, tur)
    hizala(ws)
    f = wb.create_sheet("Planogram Dışı")
    baslik(f, ["Raf", "Sıra", "Görülen", "Yüz", "Fiyat Etiketi", "Etiket Fiyatı", "Güven"], (6, 6, 40, 6, 12, 12, 9))
    for g in fazla:
        f.append([g["raf"], g["sira"], g["gorulen"], g["yuz_sayisi"], g["fiyat_etiketi"], g["etiket_fiyati"], g["guven"]])
    for b in t["bos_alanlar"]:
        f.append([b["raf"], b["sira"], "(boş alan)", "", "", "", ""])
    hizala(f)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    return {"uyumlu": uyumlu, "oran": oran}


def calistir(plan_yol: Path, foto_yollar: list[Path], cikti: Path, tur: int = 2, evet: bool = False) -> dict:
    plan = planogram_oku(plan_yol)
    fotolar = foto_hazirla(foto_yollar)
    print(f"[OK] Planogram: {len(plan)} ürün, {len({p['raf'] for p in plan})} raf · {len(fotolar)} fotoğraf")
    llm.onay_al(f"{len(fotolar)} raf fotoğrafı ve planogram listesi gönderilecek (görsel destekli model gerekir). Fotoğrafta müşteri / çalışan yüzü "
                "olmamasına dikkat edin.", evet)
    t, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(plan, fotolar, d, o), lambda x: denetle(x, plan), tur)
    satirlar, fazla = karsilastir(t, plan)
    ozet = rapor_yaz(cikti, t, satirlar, fazla, kalan, yapilan)
    return {"plan": plan, "taslak": t, "satirlar": satirlar, "fazla": fazla, "kalan": kalan, "tur": yapilan} | ozet


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    p = argparse.ArgumentParser(description="Raf fotoğrafı ve planogram listesini karşılaştırıp eksik, yanlış yerleştirilmiş ve fiyat etiketi olmayan ürünleri listeler.")
    v = BURASI / "ornek_veri"
    p.add_argument("--planogram", type=Path, default=v / "planogram.csv", help="Planogram (CSV/Excel): Raf (1 = en üst), Sıra (soldan), Barkod, Ürün Adı, Yüz Sayısı, Fiyat")
    p.add_argument("--foto", type=Path, nargs="+", default=[v / "raf_foto.jpg"], help="Raf fotoğraf(lar)ı (jpg, png …)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "planogram_uyum.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.planogram, a.foto, a.cikti, a.tur, a.evet)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] Uyum: {s['uyumlu']}/{len(s['satirlar'])} (%{s['oran']:.0f}) · planogram dışı ürün: {len(s['fazla'])}")
    for x in s["satirlar"]:
        if x["durum"] != "Yerinde" or x["sorunlar"]:
            print(f"[!] {x['ad']}: {x['durum']}; {'; '.join(x['sorunlar'])}")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
