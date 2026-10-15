"""
Şüpheli İşlem Bildirim Taslağı — Workers / Workless AI Agent
Kurumsal Uyum › Uyum Uzmanı

1. Kod işlem dökümünü ve inceleme notunu okur. Müşteri adı {MUSTERI}, karşı taraflar {KT1}, {KT2} … diye kodlanır;
   müşteri numarası modele gönderilmez, açıklamalardaki telefon / IBAN / e-posta maskelenir.
2. Kod göstergeleri hesaplar (K1, K2 …):
   - Hızlı aktarım: aynı gün giren tutarın büyük kısmının (varsayılan %90) çıkması
   - Profil uyumsuzluğu: dönem girişlerinin beyan edilen aylık gelire oranı
   - Çok sayıda farklı gönderen; çıkışların tek karşı tarafta yoğunlaşması
   - (İsteğe bağlı --esik) eşiğin hemen altında kalan işlemler (parçalama)
   Ayrıca şüphe tarihinden itibaren 10 iş günü olan son bildirim tarihini hesaplar (hafta sonu ve resmî tatiller hariç).
3. Model olay özetini ve şüphe gerekçelerini yazar; her cümle dayandığı işlem, gözlem (G) ve gösterge (K) numaralarını gösterir.
4. Kod denetler:
   - Metindeki tutar ve tarihler dökümde / hesaplarda var mı
   - Her gösterge (K) en az bir gerekçede kullanılmış mı; her gerekçe en az bir işleme ya da göstergeye dayanıyor mu
   - Kesin hüküm dili ("aklamıştır", "dolandırıcıdır", "suç geliri olduğu kesin") var mı (bildirim şüpheye dayanır)
   - Müşteriye bildirim yapılacağını ima eden ifade var mı (5549 sayılı Kanun md. 4/2: bildirim yapıldığı kimseye açıklanamaz)
   Sorun varsa modele geri gönderilir (en fazla --tur kez).
5. Kod iç taslak (gizli) belgesini oluşturur: müşteri bilgisi, son bildirim tarihi, özet tablo, olay özeti, gerekçeler.

Kullanım:
    python agent.py                                                   # örnek (kurgusal) hesap
    python agent.py --islemler dokum.xlsx --not inceleme.txt --esik 75000 --tatil 20.03.2026
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
KOLONLAR = {
    "no": ("İşlem No", "Referans", "İşlem Referansı"), "tarih": ("Tarih", "İşlem Tarihi"), "tur": ("İşlem Türü", "Tür"),
    "yon": ("Yön", "Borç/Alacak", "B/A"), "tutar": ("Tutar",), "pb": ("Para Birimi", "Döviz"), "karsi": ("Karşı Taraf", "Gönderen/Alıcı"),
    "aciklama": ("Açıklama",), "kanal": ("Kanal",),
}
NOT_ALAN = {"musteri": "Müşteri", "no": "Müşteri No", "supheli_tarih": "Şüphe Tarihi", "meslek": "Meslek", "gelir": "Beyan Edilen Aylık Gelir",
            "acilis": "Hesap Açılış", "inceleyen": "İnceleyen"}
SABIT_TATIL = [(1, 1), (4, 23), (5, 1), (5, 19), (7, 15), (8, 30), (10, 29)]
KESIN = re.compile(r"aklami\w*|aklamakta\w*|aklandigi kesin|suc geliri oldugu (kesin|acik|sabit)|teror\w* finansmani yap\w*|dolandirici(dir|sidir)|"
                   r"kesin(likle)? olarak|suclu(dur)?\b|kara para aklayici\w*|suc isle(mis|mekte)\w*|sahte hesap oldugu kesin")
IFSA = re.compile(r"musteri\w* (bildirim|haber|bilgi) ver\w*|musteri\w* bildirim yap\w*|musteri\w* bilgilendir\w*|musteri\w* (bildiril|uyarıl|uyaril)\w*|hesap sahibi\w* (bildiril|bilgilendir)\w*|"
                  r"bildirim yapildig\w* (musteri|hesap sahib)\w*")
KURUM = re.compile(r"\b(ltd|sti|as|a s|anonim|limited|sandigi|bank\w*|kurum\w*|bakanlig\w*|belediye\w*|vakf\w*|dernek\w*|sirket\w*|teknoloji|ticaret|sanayi)\b")
GELEN = {"gelen", "alacak", "a", "giris", "+"}

SEMA = {
    "type": "object",
    "properties": {
        "olay_ozeti": {"type": "array", "items": {"type": "object", "properties": {"metin": S, "dayanak": {"type": "array", "items": S}},
                                                  "required": ["metin", "dayanak"], "additionalProperties": False}},
        "supheli_gerekceler": {"type": "array", "items": {"type": "object", "properties": {"gerekce": S, "dayanak": {"type": "array", "items": S}},
                                                          "required": ["gerekce", "dayanak"], "additionalProperties": False}},
    },
    "required": ["olay_ozeti", "supheli_gerekceler"],
    "additionalProperties": False,
}


def para(v: float) -> str:
    return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def tutar_coz(x) -> float:
    if isinstance(x, (int, float)):
        return float(x)
    try:
        return abs(float(sayi_norm(str(x).replace("TL", "").replace("-", "").strip())))
    except ValueError:
        raise llm.LLMHatasi(f"Tutar okunamadı: '{x}'")


def islemleri_oku(yol: Path) -> list[dict]:
    out = []
    for r in kayitlar(yol, KOLONLAR, ("tarih", "tutar")):
        z = zaman_coz(r["tarih"])
        if not z:
            raise llm.LLMHatasi(f"{yol.name}: tarih okunamadı: '{r['tarih']}'")
        ham = str(r["tutar"])
        yon = katla(r["yon"]) or ("giden" if ham.strip().startswith("-") else "gelen")
        out.append({"no": str(r["no"] or f"I{len(out) + 1}"), "zaman": z, "tur": str(r["tur"]), "gelen": yon in GELEN,
                    "tutar": tutar_coz(r["tutar"]), "pb": str(r["pb"] or "TRY"), "karsi": str(r["karsi"]), "aciklama": str(r["aciklama"]),
                    "kanal": str(r["kanal"])})
    if not out:
        raise llm.LLMHatasi(f"{yol.name}: işlem yok.")
    return sorted(out, key=lambda x: x["zaman"])


def not_oku(yol: Path) -> dict:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: inceleme notu bulunamadı.")
    n = {a: "" for a in NOT_ALAN} | {"gozlemler": []}
    bolum = False
    for s in belge.metin_oku(yol).splitlines():
        if not s.strip():
            continue
        a, _, d = s.partition(":")
        k = next((k for k, v in NOT_ALAN.items() if katla(v) == katla(a)), None)
        if k and not s.lstrip().startswith(("-", "•", "*")):
            n[k], bolum = d.strip(), False
        elif katla(s) == "gozlemler":
            bolum = True
        elif bolum:
            n["gozlemler"].append(s.strip().lstrip("-•* ").strip())
    if not n["musteri"] or not tarih_coz(n["supheli_tarih"]):
        raise llm.LLMHatasi(f"{yol.name}: 'Müşteri:' ve 'Şüphe Tarihi: GG.AA.YYYY' satırları gerekli.")
    return n


def son_bildirim(baslangic: date, tatiller: set[date], gun: int = 10) -> date:
    d, say = baslangic, 0
    while say < gun:
        d += timedelta(days=1)
        if d.weekday() < 5 and d not in tatiller and (d.month, d.day) not in SABIT_TATIL:
            say += 1
    return d


def kodla(islemler: list[dict], n: dict) -> dict:
    harita = {n["musteri"]: "{MUSTERI}"}
    for x in islemler:
        if x["karsi"] and x["karsi"] not in harita:
            harita[x["karsi"]] = "{KT%d}" % (len(harita))
    return harita


def gizle(metin: str, harita: dict) -> str:
    for ad, kod in sorted(harita.items(), key=lambda k: -len(k[0])):
        metin = re.sub(re.escape(ad), kod, metin, flags=re.I)
        if len(ad.split()) in (2, 3) and not KURUM.search(katla(ad)):  # kişi adı: yalnız adıyla ya da soyadıyla geçen yerler
            for parca in ad.split():
                if len(parca) > 3 and parca[0].isupper():
                    metin = re.sub(rf"\b{re.escape(parca)}\b", kod, metin)
    return llm.maskele(metin)


def gostergeler(islemler: list[dict], n: dict, oran: float, esik: float | None, harita: dict) -> list[dict]:
    k = []
    gun = defaultdict(lambda: {"giris": 0.0, "cikis": 0.0, "no": []})
    for x in islemler:
        g = gun[x["zaman"].date()]
        g["giris" if x["gelen"] else "cikis"] += x["tutar"]
        g["no"].append(x["no"])
    hizli = [(d, g) for d, g in sorted(gun.items()) if g["giris"] > 0 and g["cikis"] >= g["giris"] * oran]
    if hizli:
        k.append({"ad": "Hızlı aktarım", "no": [i for _, g in hizli for i in g["no"]],
                  "metin": "Aynı gün giren tutarın %" + f"{oran * 100:.0f}" + "'ından fazlası çıkmış: " +
                           "; ".join(f"{d:%d.%m.%Y} giriş {para(g['giris'])} TL, çıkış {para(g['cikis'])} TL" for d, g in hizli)})
    giris = sum(x["tutar"] for x in islemler if x["gelen"])
    gelir = tutar_coz(n["gelir"]) if str(n["gelir"]).strip() else 0
    if gelir:
        bas, son = islemler[0]["zaman"].date(), islemler[-1]["zaman"].date()
        k.append({"ad": "Profil uyumsuzluğu", "no": [x["no"] for x in islemler if x["gelen"]],
                  "metin": f"{bas:%d.%m.%Y}-{son:%d.%m.%Y} döneminde toplam giriş {para(giris)} TL; beyan edilen aylık gelirin ({para(gelir)} TL) "
                           + f"{giris / gelir:.1f}".replace(".", ",") + " katı"})
    gonderen = {x["karsi"] for x in islemler if x["gelen"] and x["karsi"] and katla(x["tur"]) in {"eft", "havale", "fast", "swift"}}
    if len(gonderen) >= 5:
        k.append({"ad": "Çok sayıda farklı gönderen", "no": [x["no"] for x in islemler if x["gelen"] and x["karsi"] in gonderen],
                  "metin": f"{len(gonderen)} farklı kişi / kurumdan EFT / havale girişi"})
    cikis = [x for x in islemler if not x["gelen"]]
    if cikis:
        top = defaultdict(float)
        for x in cikis:
            top[x["karsi"]] += x["tutar"]
        kim, tutar = max(top.items(), key=lambda t: t[1])
        toplam = sum(top.values())
        if kim and tutar >= toplam * 0.7 and sum(1 for x in cikis if x["karsi"] == kim) >= 2:
            k.append({"ad": "Çıkışlarda yoğunlaşma", "no": [x["no"] for x in cikis if x["karsi"] == kim],
                      "metin": f"Toplam {para(toplam)} TL çıkışın {para(tutar)} TL'si {harita.get(kim, kim)} hesabına "
                               f"({sum(1 for x in cikis if x['karsi'] == kim)} işlem)"})
    if esik:
        alt = [x for x in islemler if esik * 0.9 <= x["tutar"] < esik]
        if len(alt) >= 2:
            k.append({"ad": "Eşik altı (parçalama)", "no": [x["no"] for x in alt],
                      "metin": f"{len(alt)} işlem {para(esik)} TL eşiğinin %90'ı ile eşik arasında"})
    for i, x in enumerate(k, 1):
        x["id"] = f"K{i}"
    return k


def ozet(islemler: list[dict]) -> dict:
    giris = sum(x["tutar"] for x in islemler if x["gelen"])
    cikis = sum(x["tutar"] for x in islemler if not x["gelen"])
    return {"adet": len(islemler), "giris": giris, "cikis": cikis, "bas": islemler[0]["zaman"], "son": islemler[-1]["zaman"]}


def kimlik_haritasi(islemler, n, k) -> dict:
    return {x["no"]: x for x in islemler} | {f"G{i}": g for i, g in enumerate(n["gozlemler"], 1)} | {x["id"]: x for x in k}


def izinli_sayilar(islemler, n, k, oz, harita) -> set:
    parcalar = [f"{x['zaman']:%d.%m.%Y %H:%M} {para(x['tutar'])}" for x in islemler] + [gizle(g, harita) for g in n["gozlemler"]]
    # Modele verilen müşteri profilindeki hesap açılış tarihi de izinli
    parcalar += [x["metin"] for x in k] + [f"{oz['adet']} {para(oz['giris'])} {para(oz['cikis'])} {n['gelir']} {n['supheli_tarih']} {n['acilis'] or ''}"]
    s = sayilar(" ".join(parcalar))
    return s | {re.sub(r"\D", "", m) for m in re.findall(r"\{KT\d+\}", " ".join(parcalar))}


def denetle(t: dict, islemler, n, k, oz, harita) -> list[str]:
    sorun = []
    kim = kimlik_haritasi(islemler, n, k)
    izinli = izinli_sayilar(islemler, n, k, oz, harita)
    kodlar = set(harita.values())
    kullanilan = set()
    parcalar = [(f"Özet {i}", p["metin"], p["dayanak"]) for i, p in enumerate(t["olay_ozeti"], 1)]
    parcalar += [(f"Gerekçe {i}", g["gerekce"], g["dayanak"]) for i, g in enumerate(t["supheli_gerekceler"], 1)]
    for ad, metin, dayanak in parcalar:
        gec = [d for d in dayanak if d not in kim]
        if not dayanak or gec:
            sorun.append(f"{ad}: dayanak yok ya da geçersiz {gec}")
        if ad.startswith("Gerekçe"):
            kullanilan |= set(dayanak)
            if dayanak and all(d.startswith("G") and d[1:].isdigit() for d in dayanak):
                sorun.append(f"{ad}: yalnız gözleme dayanıyor; en az bir işlem ya da gösterge (K) numarası ekleyin")
        temiz = re.sub(r"\{(MUSTERI|KT\d+)\}", " ", metin)
        fazla = sorted(sayilar(temiz) - izinli)
        if fazla:
            sorun.append(f"{ad}: dökümde / hesaplarda olmayan sayı: {', '.join(fazla)}")
        for yer in set(re.findall(r"\{[A-Z]+\d*\}", metin)) - kodlar:
            sorun.append(f"{ad}: tanımsız yer tutucu {yer}")
        km = katla(metin)
        if KESIN.search(km):
            sorun.append(f"{ad}: kesin hüküm dili ('{KESIN.search(km).group(0)}'); bildirim şüpheye dayanır, suç tespiti yapılmaz")
        if IFSA.search(km):
            sorun.append(f"{ad}: müşteriye bildirim yapılacağını ima ediyor; 5549 sayılı Kanun md. 4/2 gereği bildirim kimseye açıklanamaz")
    for x in k:
        if x["id"] not in kullanilan:
            sorun.append(f"{x['id']}: gösterge ({x['ad']}) hiçbir şüphe gerekçesinde kullanılmamış")
    if not t["supheli_gerekceler"]:
        sorun.append("Genel: şüphe gerekçesi yok")
    return sorun


def uret(islemler, n, k, oz, harita, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    satir = "\n".join(f"{x['no']} | {x['zaman']:%d.%m.%Y %H:%M} | {x['tur']} | {'Gelen' if x['gelen'] else 'Giden'} | {para(x['tutar'])} {x['pb']} | "
                      f"{harita.get(x['karsi'], '-') if x['karsi'] else '-'} | {gizle(x['aciklama'], harita)} | {x['kanal']}" for x in islemler)
    g = "\n".join(f"G{i} | {gizle(s, harita)}" for i, s in enumerate(n["gozlemler"], 1)) or "(yok)"
    kk = "\n".join(f"{x['id']} | {x['ad']} | {x['metin']}" for x in k) or "(kodun bulduğu gösterge yok)"
    mesaj = (f"<musteri_profili>Müşteri: {{MUSTERI}} · Meslek: {gizle(n['meslek'], harita) or '-'} · Beyan edilen aylık gelir: {n['gelir'] or '-'} TL · "
             f"Hesap açılış: {n['acilis'] or '-'}</musteri_profili>\n"
             f"<ozet>{oz['adet']} işlem · toplam giriş {para(oz['giris'])} TL · toplam çıkış {para(oz['cikis'])} TL · "
             f"{oz['bas']:%d.%m.%Y}-{oz['son']:%d.%m.%Y}</ozet>\n<islemler>\n{satir}\n</islemler>\n<gozlemler>\n{g}\n</gozlemler>\n"
             f"<gostergeler>\n{kk}\n</gostergeler>" + duzeltme_blogu(duzeltme, onceki))
    return llm.json_iste(sistem, mesaj, SEMA)


def geri_ac(metin: str, harita: dict) -> str:
    for ad, kod in harita.items():
        metin = metin.replace(kod, ad)
    return metin


def belge_metni(t, islemler, n, k, oz, harita, son) -> str:
    s = ["ŞÜPHELİ İŞLEM BİLDİRİMİ — İÇ TASLAK (GİZLİ)", "",
         "Bu belge MASAK'a yapılacak bildirime hazırlık amaçlı iç taslaktır. 5549 sayılı Kanun md. 4/2 uyarınca bildirimde bulunulduğu, "
         "yükümlülük denetimi elemanları ve yargılama sırasında mahkemeler dışında, işlemin tarafları dâhil hiç kimseye açıklanamaz.", "",
         f"Müşteri: {n['musteri']} · Müşteri No: {n['no'] or '-'}", f"Meslek / beyan edilen aylık gelir: {n['meslek'] or '-'} / {n['gelir'] or '-'} TL",
         f"Şüphenin oluştuğu tarih: {n['supheli_tarih']} · Son bildirim tarihi (10 iş günü): {son:%d.%m.%Y} (gecikmesinde sakınca varsa derhal)",
         f"İncelenen dönem: {oz['bas']:%d.%m.%Y}-{oz['son']:%d.%m.%Y} · {oz['adet']} işlem · giriş {para(oz['giris'])} TL · çıkış {para(oz['cikis'])} TL",
         "", "GÖSTERGELER (kod hesabı)"] + [f"  {x['id']}. {x['ad']}: {geri_ac(x['metin'], harita)}" for x in k]
    s += ["", "OLAY ÖZETİ"] + [geri_ac(p["metin"], harita) + f" [{', '.join(p['dayanak'])}]" for p in t["olay_ozeti"]]
    s += ["", "ŞÜPHE GEREKÇELERİ"] + [f"  {i}. {geri_ac(g['gerekce'], harita)} [{', '.join(g['dayanak'])}]" for i, g in enumerate(t["supheli_gerekceler"], 1)]
    s += ["", "Bildirim kararı ve MASAK Online sistemine giriş uyum görevlisinin değerlendirmesiyle yapılır."]
    return "\n".join(s)


def rapor_yaz(cikti: Path, t, islemler, n, k, oz, harita, son, kalan, tur, metin) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Taslak"
    baslik(ws, ["Bölüm", "Metin", "Dayanak"], (14, 100, 30))
    for i, p in enumerate(t["olay_ozeti"], 1):
        ws.append([f"Özet {i}", geri_ac(p["metin"], harita), ", ".join(p["dayanak"])])
        boya(ws, (2,), AI_DOLGU)
    for i, g in enumerate(t["supheli_gerekceler"], 1):
        ws.append([f"Gerekçe {i}", geri_ac(g["gerekce"], harita), ", ".join(g["dayanak"])])
        boya(ws, (2,), AI_DOLGU)
    son_not(ws, f"GİZLİ. Son bildirim tarihi: {son:%d.%m.%Y}. Taslaktır; bildirim kararı uyum görevlisine aittir. Bildirim yapıldığı kimseye açıklanamaz "
                "(5549 md. 4/2).", kalan, tur)
    hizala(ws)
    w = wb.create_sheet("İşlemler")
    baslik(w, ["İşlem No", "Tarih", "Tür", "Yön", "Tutar", "PB", "Karşı Taraf", "Açıklama", "Kanal", "Göstergeler"], (12, 17, 14, 8, 14, 6, 30, 20, 10, 20))
    for x in islemler:
        g = [y["id"] for y in k if x["no"] in y["no"]]
        w.append([x["no"], x["zaman"], x["tur"], "Gelen" if x["gelen"] else "Giden", x["tutar"], x["pb"], x["karsi"], x["aciklama"], x["kanal"], ", ".join(g)])
        w.cell(w.max_row, 2).number_format = "DD.MM.YYYY HH:MM"
        w.cell(w.max_row, 5).number_format = "#,##0.00"
        if g:
            boya(w, (10,), SARI)
    hizala(w)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    cikti.with_suffix(".txt").write_text(metin, encoding="utf-8")


def calistir(islem_yol: Path, not_yol: Path, cikti: Path, oran: float = 0.9, esik: float | None = None, tatiller: set | None = None,
             tur: int = 2, evet: bool = False) -> dict:
    islemler = islemleri_oku(islem_yol)
    n = not_oku(not_yol)
    harita = kodla(islemler, n)
    k = gostergeler(islemler, n, oran, esik, harita)
    oz = ozet(islemler)
    son = son_bildirim(tarih_coz(n["supheli_tarih"]), tatiller or set())
    print(f"[OK] {oz['adet']} işlem · {len(k)} gösterge · son bildirim {son:%d.%m.%Y}")
    llm.onay_al(f"{oz['adet']} işlem, {len(n['gozlemler'])} gözlem ve göstergeler gönderilecek; müşteri ve karşı taraf adları kodlanır, müşteri no gönderilmez.", evet)
    t, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(islemler, n, k, oz, harita, d, o), lambda x: denetle(x, islemler, n, k, oz, harita), tur)
    metin = belge_metni(t, islemler, n, k, oz, harita, son)
    rapor_yaz(cikti, t, islemler, n, k, oz, harita, son, kalan, yapilan, metin)
    return {"gostergeler": k, "son_bildirim": son, "taslak": t, "kalan": kalan, "tur": yapilan, "metin": metin, "harita": harita}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    p = argparse.ArgumentParser(description="İnceleme notları ve işlem dökümünden MASAK şüpheli işlem bildirimi için olay özeti ve şüphe gerekçesi taslağı hazırlar.")
    v = BURASI / "ornek_veri"
    p.add_argument("--islemler", type=Path, default=v / "islemler.csv", help="İşlem dökümü (CSV/Excel): İşlem No, Tarih, İşlem Türü, Yön, Tutar, Para Birimi, Karşı Taraf, Açıklama, Kanal")
    p.add_argument("--not", dest="not_", type=Path, default=v / "inceleme_notu.txt",
                   help="İnceleme notu (.txt/.docx): Müşteri, Müşteri No, Şüphe Tarihi, Meslek, Beyan Edilen Aylık Gelir satırları; 'Gözlemler' bölümü")
    p.add_argument("--oran", type=float, default=0.9, help="Hızlı aktarım için aynı gün çıkış / giriş oranı (varsayılan 0,9)")
    p.add_argument("--esik", type=float, help="Parçalama kontrolü için tutar eşiği (TL); verilmezse kontrol yapılmaz")
    p.add_argument("--tatil", nargs="*", default=[], help="Ek resmî tatil günleri (GG.AA.YYYY), ör. dini bayramlar")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "supheli_islem_taslagi.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        tatiller = set()
        for x in a.tatil:
            if not tarih_coz(x):
                raise llm.LLMHatasi(f"Tatil tarihi okunamadı: '{x}' (GG.AA.YYYY)")
            tatiller.add(tarih_coz(x))
        s = calistir(a.islemler, a.not_, a.cikti, a.oran, a.esik, tatiller, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    for x in s["gostergeler"]:
        print(f"[i] {x['id']} {x['ad']}")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Taslak: {a.cikti.with_suffix('.txt').resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
