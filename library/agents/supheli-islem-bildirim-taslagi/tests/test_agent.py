"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

V = KLASOR / "ornek_veri"
ISL, NOT = V / "islemler.csv", V / "inceleme_notu.txt"
MESAJLAR = []


def p(metin, *d):
    return {"metin": metin, "dayanak": list(d)}


def g(metin, *d):
    return {"gerekce": metin, "dayanak": list(d)}


def taslak(dogru=True) -> dict:
    t = {"olay_ozeti": [
        p("{MUSTERI}, emekli öğretmen olup beyan edilen aylık geliri 32.000 TL'dir. Hesap uzun süredir yalnız emekli maaşı ve küçük tutarlı fatura "
          "ödemeleri için kullanılmaktadır.", "G1", "TX1012"),
        p("01.10.2026-08.10.2026 döneminde hesaba 7 farklı kişiden 45.000,00 TL ile 49.500,00 TL arasında EFT girişleri olmuş, toplanan tutarlar aynı "
          "gün {KT3} hesabına gönderilmiştir (ör. 03.10.2026 tarihinde 140.000,00 TL).", "TX1001", "TX1007", "K1", "K3", "K4"),
        p("Müşteri şubede işlemler hakkında tanıdıklara aracılık ettiğini belirtmiş, belge sunmamıştır.", "G6")],
        "supheli_gerekceler": [
            g("Aynı gün giren tutarların neredeyse tamamının hesaptan çıkması, hesabın aracı (geçiş) hesap olarak kullanıldığı şüphesini doğurmaktadır.", "K1", "TX1003"),
            g("Dönem girişleri beyan edilen aylık gelirin 11,4 katıdır; müşterinin profiliyle uyumlu görünmemektedir.", "K2"),
            g("Daha önce işlemi olmayan 7 farklı kişiden genel açıklamalı EFT'ler alınması, ekonomik amacı anlaşılamayan bir işlem izlenimi vermektedir.", "K3", "G3"),
            g("Çıkışların {KT3} hesabında yoğunlaşması ve yeni cihazdan oturum açılması hesabın üçüncü kişilerce kullanıldığı şüphesini doğurmaktadır.", "K4", "G5")]}
    if not dogru:
        t["olay_ozeti"][1]["metin"] = t["olay_ozeti"][1]["metin"].replace("140.000,00", "145.000,00")
        t["supheli_gerekceler"] = t["supheli_gerekceler"][:2]
        t["supheli_gerekceler"][1] = g("Müşteri, suç gelirini kesinlikle aklamıştır; müşteriye bildirim yapılacağı bilgisi verilmelidir.", "G6")
    return t


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return taslak("<duzeltme>" in kullanici)


def veri():
    i, n = agent.islemleri_oku(ISL), agent.not_oku(NOT)
    h = agent.kodla(i, n)
    return i, n, h, agent.gostergeler(i, n, 0.9, None, h), agent.ozet(i)


class KodTesti(unittest.TestCase):
    def test_hesap(self):
        i, n, h, k, oz = veri()
        self.assertEqual((oz["adet"], oz["giris"], oz["cikis"]), (12, 365500.0, 329000.0))
        self.assertEqual([x["ad"] for x in k], ["Hızlı aktarım", "Profil uyumsuzluğu", "Çok sayıda farklı gönderen", "Çıkışlarda yoğunlaşma"])
        self.assertIn("11,4 katı", k[1]["metin"])
        self.assertEqual(k[3]["no"], ["TX1003", "TX1007", "TX1011"])
        self.assertEqual(len(agent.gostergeler(i, n, 0.9, 50000, h)[4]["no"]), 7)
        self.assertEqual(agent.son_bildirim(date(2026, 10, 9), set()), date(2026, 10, 23))
        self.assertEqual(agent.son_bildirim(date(2026, 10, 22), set()), date(2026, 11, 6))  # 29 Ekim atlanır
        self.assertEqual(agent.gizle("Selin Hanım ile Hasan Er görüştü, tel 0532 111 22 33", h), "{KT2} Hanım ile {MUSTERI} görüştü, tel [TELEFON]")
        self.assertEqual(agent.gizle("Emekli öğretmen", h), "Emekli öğretmen")

    def test_denetim(self):
        i, n, h, k, oz = veri()
        self.assertEqual(agent.denetle(taslak(), i, n, k, oz, h), [])
        s = agent.denetle(taslak(False), i, n, k, oz, h)
        for parca in ("145000", "yalnız gözleme", "kesin hüküm", "5549", "K2:", "K3:", "K4:"):
            self.assertTrue(any(parca in x for x in s), (parca, s))
        self.assertEqual(len(s), 7, s)


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "s.xlsx"
            s = agent.calistir(ISL, NOT, cikti, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            for gizli in ("Hasan", "Ahmet", "Yılmaz", "Deneme Teknoloji", "77001234", "Zeynep"):
                self.assertNotIn(gizli, "".join(MESAJLAR))
            m = cikti.with_suffix(".txt").read_text(encoding="utf-8")
            self.assertTrue(m.startswith("ŞÜPHELİ İŞLEM BİLDİRİMİ — İÇ TASLAK (GİZLİ)"))
            self.assertIn("Müşteri: Hasan Er · Müşteri No: 77001234", m)
            self.assertIn("Son bildirim tarihi (10 iş günü): 23.10.2026", m)
            self.assertIn("Deneme Teknoloji Ltd. Şti. hesabına", m)
            ws = load_workbook(cikti)["İşlemler"]
            self.assertEqual((ws["A4"].value, ws["J4"].value), ("TX1003", "K1, K4"))

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(ISL, NOT, Path(tmp) / "s.xlsx", tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 7)
            self.assertEqual(agent.main(["--evet", "--esik", "50000", "--tatil", "20.10.2026", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--evet", "--tatil", "bayram"]), 1)


if __name__ == "__main__":
    unittest.main()
