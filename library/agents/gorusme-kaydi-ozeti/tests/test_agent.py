"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import re
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri" / "gorusmeler.csv"
MESAJLAR = []
YANIT = {
    "G-5001": ("beklemede", [("Kargo firmasından dönüş al", "Temsilci", "yarın")], [("Yarın bilgi verilecek", "yarın")]),
    "G-5002": ("beklemede", [("Fatura yeniden düzenlenecek", "Muhasebe", "3 iş günü içinde")], [("Yeni fatura e-postayla gönderilecek", "3 iş günü içinde")]),
    "G-5003": ("cozuldu", [], []),
    "G-5004": ("eskalasyon", [("Para iadesi kontrolü", "Finans", "")], [("Müşteri aranacak", "bu hafta içinde")]),
    "G-5005": ("eskalasyon", [("Kargo konusu", "", "")], []),
}


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    idler = re.findall(r'<gorusme id="([^"]+)"', kullanici)
    return {"gorusmeler": [{"id": i, "kategori": "Diğer", "talep": "t", "yapilan_islem": "i", "durum": YANIT[i][0],
                            "aksiyonlar": [{"is": a, "sorumlu": b, "sure_ifadesi": c} for a, b, c in YANIT[i][1]],
                            "taahhutler": [{"soz": a, "sure_ifadesi": c} for a, c in YANIT[i][2]],
                            "duygu": "notr", "risk": "Hakem heyeti" if i == "G-5004" else ""} for i in idler]}


class TarihTesti(unittest.TestCase):
    def test_is_gunu_ve_29_ekim(self):
        # Salı 27.10.2026 + 3 iş günü: Çar 28 (1), Per 29 Cumhuriyet Bayramı atlanır, Cum 30 (2), Pzt 2 Kasım (3)
        self.assertEqual(agent.sure_coz("3 iş günü içinde", date(2026, 10, 27), set()), date(2026, 11, 2))
        self.assertEqual(agent.sure_coz("üç iş günü", date(2026, 10, 27), set()), date(2026, 11, 2))

    def test_diger_ifadeler(self):
        b = date(2026, 10, 28)                                            # Çarşamba
        self.assertEqual(agent.sure_coz("yarın", b, set()), date(2026, 10, 29))
        self.assertEqual(agent.sure_coz("bu hafta içinde", b, set()), date(2026, 10, 30))
        self.assertEqual(agent.sure_coz("haftaya", b, set()), date(2026, 11, 6))
        self.assertEqual(agent.sure_coz("5 gün içinde", b, set()), date(2026, 11, 2))
        self.assertEqual(agent.sure_coz("15.11.2026 tarihine kadar", b, set()), date(2026, 11, 15))
        self.assertEqual(agent.sure_coz("ay sonuna kadar", b, set()), date(2026, 10, 30))   # 31 Ekim cumartesi
        self.assertIsNone(agent.sure_coz("en kısa sürede", b, set()))

    def test_ekstra_tatil(self):
        self.assertEqual(agent.sure_coz("1 iş günü", date(2026, 10, 27), {date(2026, 10, 28)}), date(2026, 10, 30))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_uctan_uca(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "g.xlsx"
            s = agent.calistir(ORNEK, cikti, bugun=date(2026, 10, 29), evet=True)
            m = MESAJLAR[0]
            for x in ("Ahmet Örnek", "Selin Deneme", "0533 444 55 66"):
                self.assertNotIn(x, m)
            self.assertIn('onceki="G-5001"', m)                           # tekrar arayan müşteri bağlandı
            self.assertEqual(s["sonuc"]["G-5002"]["aksiyonlar"][0]["son_tarih"], date(2026, 11, 2))
            self.assertEqual(s["sonuc"]["G-5004"]["taahhutler"][0]["son_tarih"], date(2026, 10, 30))
            self.assertIn("tarihi olmayan aksiyon: Para iadesi kontrolü", s["kontroller"]["G-5004"])
            self.assertIn("sorumlusu olmayan aksiyon: Kargo konusu", s["kontroller"]["G-5005"])
            ws = load_workbook(cikti)["Takip Listesi"]
            satirlar = list(ws.iter_rows(min_row=2, values_only=True))
            self.assertEqual(satirlar[0][1], "GECİKMİŞ")                   # 28.10 yarın sözü, bugün 29.10
            self.assertEqual(satirlar[-1][1], "Tarihsiz")


if __name__ == "__main__":
    unittest.main()
