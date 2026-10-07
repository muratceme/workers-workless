"""
Görüşme Kaydı Özeti — Workers / Workless AI Agent
Müşteri Hizmetleri ve Çağrı Merkezi › Müşteri Temsilcisi

1. Görüşme notları okunur; aynı müşterinin tekrar eden görüşmeleri kodla bağlanır. Müşteri adları takma adla
   (K1, K2...) değiştirilir; telefon, e-posta, IBAN ve TCKN maskelenir.
2. Model her notu standart kayda çevirir: kategori, talep, yapılan işlem, durum (çözüldü / beklemede /
   eskalasyon), bekleyen aksiyonlar (sorumlu + süre ifadesi), müşteriye verilen sözler, duygu ve risk.
3. Kod süre ifadelerini ("yarın", "3 iş günü içinde", "bu hafta içinde", "15.11.2026") görüşme tarihine göre
   gerçek son tarihe çevirir. İş günü hesabında hafta sonu ve sabit tarihli resmî tatiller atlanır (dini
   bayramlar --tatiller dosyasıyla eklenir). Bekleyen ama sorumlusu ya da tarihi olmayan kayıtlar işaretlenir.
4. Çıktı: CRM'e aktarılabilir kayıt listesi + son tarihe göre sıralı takip listesi (Excel).

Kullanım:
    python agent.py                                           # örnek görüşmelerle dener
    python agent.py --girdi gorusmeler.xlsx --tatiller dini_bayramlar_2026.txt
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import OrderedDict
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import llm

BURASI = Path(__file__).resolve().parent
PAKET = 20
KATEGORILER = ["Sipariş ve kargo", "İade ve değişim", "Fatura", "Ödeme ve para iadesi", "Ürün bilgisi", "Şikâyet", "Teknik destek",
               "Hesap ve üyelik", "KVKK", "Diğer"]
# 2429 sayılı Kanun: sabit tarihli resmî tatiller (dini bayramlar her yıl değişir → --tatiller)
SABIT_TATILLER = {(1, 1), (4, 23), (5, 1), (5, 19), (7, 15), (8, 30), (10, 29)}
SAYILAR = {"bir": 1, "iki": 2, "üç": 3, "dört": 4, "beş": 5, "altı": 6, "yedi": 7, "sekiz": 8, "dokuz": 9, "on": 10}


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def zaman(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    if isinstance(x, date):
        return datetime(x.year, x.month, x.day)
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%d.%m.%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(x).strip(), f)
        except ValueError:
            pass
    return None


# ----------------------------------------------------------------------------
# Süre ifadesi → tarih (deterministik)
# ----------------------------------------------------------------------------

def is_gunu_mu(t: date, ekstra: set[date]) -> bool:
    return t.weekday() < 5 and (t.month, t.day) not in SABIT_TATILLER and t not in ekstra


def is_gunu_ekle(t: date, n: int, ekstra: set[date]) -> date:
    while n > 0:
        t += timedelta(days=1)
        if is_gunu_mu(t, ekstra):
            n -= 1
    return t


def son_is_gunu(t: date, ekstra: set[date]) -> date:
    while not is_gunu_mu(t, ekstra):
        t -= timedelta(days=1)
    return t


def sure_coz(ifade: str, baslangic: date, ekstra: set[date]) -> date | None:
    s = kucuk(ifade)
    if not s:
        return None
    m = re.search(r"\b(\d{1,2})[./](\d{1,2})(?:[./](\d{4}))?\b", s)
    if m:
        y = int(m.group(3)) if m.group(3) else baslangic.year
        try:
            t = date(y, int(m.group(2)), int(m.group(1)))
        except ValueError:
            return None
        return t if m.group(3) or t >= baslangic else date(y + 1, t.month, t.day)
    for k, v in SAYILAR.items():
        s = re.sub(rf"\b{k}\b", str(v), s)
    if "bugün" in s:
        return baslangic
    if "yarın" in s:
        return baslangic + timedelta(days=1)
    m = re.search(r"(\d+)\s*iş\s*gün", s)
    if m:
        return is_gunu_ekle(baslangic, int(m.group(1)), ekstra)
    m = re.search(r"(\d+)\s*gün", s)
    if m:
        return baslangic + timedelta(days=int(m.group(1)))
    m = re.search(r"(\d+)\s*hafta", s)
    if m:
        return baslangic + timedelta(days=7 * int(m.group(1)))
    cuma = baslangic + timedelta(days=4 - baslangic.weekday())
    if re.search(r"bu hafta", s):
        return son_is_gunu(cuma if cuma >= baslangic else baslangic, ekstra)
    if re.search(r"(haftaya|gelecek hafta|önümüzdeki hafta)", s):
        return son_is_gunu(cuma + timedelta(days=7), ekstra)
    if re.search(r"(ay sonu|bu ay)", s):
        sonraki = date(baslangic.year + (baslangic.month == 12), baslangic.month % 12 + 1, 1)
        return son_is_gunu(sonraki - timedelta(days=1), ekstra)
    return None


# ----------------------------------------------------------------------------
# Okuma
# ----------------------------------------------------------------------------

def gorusmeleri_oku(yol: Path) -> list[dict]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    s = [r for r in s if any(c not in (None, "") for c in r)]
    b = [kucuk(x) for x in s[0]]
    bul = lambda *a: next((i for i, x in enumerate(b) if x in a), None)  # noqa: E731
    k = {"id": bul("görüşme no", "kayıt no", "id", "no"), "tarih": bul("tarih", "görüşme tarihi"), "kanal": bul("kanal"),
         "temsilci": bul("temsilci", "agent"), "musteri": bul("müşteri", "müşteri no", "arayan"), "not": bul("not", "görüşme notu", "açıklama")}
    if k["not"] is None or k["tarih"] is None:
        raise SystemExit(f"Tarih ve Not sütunları gerekli. Başlıklar: {s[0]}")
    al = lambda r, a: r[k[a]] if k[a] is not None and k[a] < len(r) else None  # noqa: E731
    return [{"id": str(al(r, "id") or f"G{i}"), "zaman": zaman(al(r, "tarih")), "kanal": str(al(r, "kanal") or ""),
             "temsilci": str(al(r, "temsilci") or ""), "musteri": str(al(r, "musteri") or "").strip(), "not": str(al(r, "not") or "").strip()}
            for i, r in enumerate(s[1:], 1) if str(al(r, "not") or "").strip()]


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

_SURELI = {"type": "object", "properties": {"is": {"type": "string"}, "sorumlu": {"type": "string"}, "sure_ifadesi": {"type": "string"}},
           "required": ["is", "sorumlu", "sure_ifadesi"], "additionalProperties": False}
SEMA = {
    "type": "object",
    "properties": {"gorusmeler": {"type": "array", "items": {
        "type": "object",
        "properties": {
            "id": {"type": "string"},
            "kategori": {"type": "string", "enum": KATEGORILER},
            "talep": {"type": "string"},
            "yapilan_islem": {"type": "string"},
            "durum": {"type": "string", "enum": ["cozuldu", "beklemede", "eskalasyon"]},
            "aksiyonlar": {"type": "array", "items": _SURELI},
            "taahhutler": {"type": "array", "items": {"type": "object", "properties": {"soz": {"type": "string"}, "sure_ifadesi": {"type": "string"}},
                                                       "required": ["soz", "sure_ifadesi"], "additionalProperties": False}},
            "duygu": {"type": "string", "enum": ["memnun", "notr", "memnuniyetsiz", "ofkeli"]},
            "risk": {"type": "string"},
        },
        "required": ["id", "kategori", "talep", "yapilan_islem", "durum", "aksiyonlar", "taahhutler", "duygu", "risk"],
        "additionalProperties": False}}},
    "required": ["gorusmeler"],
    "additionalProperties": False,
}


def ozetle(gorusmeler: list[dict], harita: dict[str, str]) -> dict[str, dict]:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    sonuc = {}
    for i in range(0, len(gorusmeler), PAKET):
        parca = gorusmeler[i:i + PAKET]
        satirlar = []
        for g in parca:
            not_ = llm.maskele(g["not"])
            if g["musteri"]:
                not_ = re.sub(re.escape(g["musteri"]), harita[g["musteri"]], not_, flags=re.I)
            satirlar.append(f'<gorusme id="{g["id"]}" tarih="{g["zaman"]:%d.%m.%Y %H:%M}" kanal="{g["kanal"]}" '
                            f'musteri="{harita.get(g["musteri"], "-")}" onceki="{", ".join(g["onceki"]) or "-"}">\n{not_}\n</gorusme>')
        mesaj = "\n".join(["<kategoriler>" + "; ".join(KATEGORILER) + "</kategoriler>", "<gorusmeler>", *satirlar, "</gorusmeler>"])
        yanit = llm.json_iste(sistem, mesaj, SEMA)
        gecerli = {g["id"] for g in parca}
        for s in yanit.get("gorusmeler", []):
            if s.get("id") in gecerli:
                sonuc[s["id"]] = s
    return sonuc


def kontrol_et(g: dict, s: dict, ekstra: set[date]) -> list[str]:
    notlar = []
    gun = g["zaman"].date()
    for a in s["aksiyonlar"] + [{"is": t["soz"], "sorumlu": "Temsilci", "sure_ifadesi": t["sure_ifadesi"]} for t in s["taahhutler"]]:
        a["son_tarih"] = sure_coz(a["sure_ifadesi"], gun, ekstra)
        if a["sure_ifadesi"] and a["son_tarih"] is None:
            notlar.append(f"süre ifadesi tarihe çevrilemedi: '{a['sure_ifadesi']}'")
    for t in s["taahhutler"]:
        t["son_tarih"] = sure_coz(t["sure_ifadesi"], gun, ekstra)
    if s["durum"] != "cozuldu":
        if not s["aksiyonlar"]:
            notlar.append("beklemede/eskalasyon ama bekleyen aksiyon yazılmamış")
        for a in s["aksiyonlar"]:
            if not a["sorumlu"].strip():
                notlar.append(f"sorumlusu olmayan aksiyon: {a['is'][:40]}")
            if not a["sure_ifadesi"].strip():
                notlar.append(f"tarihi olmayan aksiyon: {a['is'][:40]}")
    return notlar


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
UST = Alignment(vertical="top", wrap_text=True)
DURUM = {"cozuldu": "Çözüldü", "beklemede": "Beklemede", "eskalasyon": "Eskalasyon"}
DUYGU = {"memnun": "Memnun", "notr": "Nötr", "memnuniyetsiz": "Memnuniyetsiz", "ofkeli": "Öfkeli"}


def _baslik(ws, b, g):
    ws.append(b)
    for c in ws[ws.max_row]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI
    for j, w in enumerate(g, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "B2"


def rapor_yaz(cikti: Path, gorusmeler: list[dict], sonuc: dict, kontroller: dict, bugun: date) -> None:
    wb = Workbook()
    t = wb.active
    t.title = "Takip Listesi"
    _baslik(t, ["Son Tarih", "Durum", "Görüşme No", "Müşteri", "İş / Söz", "Sorumlu", "Süre İfadesi", "Tür", "Tamamlandı"],
            (12, 12, 11, 18, 60, 18, 20, 10, 11))
    satirlar = []
    for g in gorusmeler:
        s = sonuc.get(g["id"])
        if not s:
            continue
        for a in s["aksiyonlar"]:
            satirlar.append((a.get("son_tarih"), g, a["is"], a["sorumlu"], a["sure_ifadesi"], "Aksiyon"))
        for x in s["taahhutler"]:
            satirlar.append((x.get("son_tarih"), g, x["soz"], "Temsilci", x["sure_ifadesi"], "Söz"))
    for son, g, is_, sorumlu, ifade, tur in sorted(satirlar, key=lambda x: (x[0] is None, x[0] or date.max)):
        durum = "Tarihsiz" if son is None else "GECİKMİŞ" if son < bugun else "Bugün" if son == bugun else "Açık"
        t.append([son, durum, g["id"], g["musteri"], is_, sorumlu, ifade, tur, ""])
        t.cell(t.max_row, 1).number_format = "DD.MM.YYYY"
        t.cell(t.max_row, 2).fill = KIRMIZI if durum in ("GECİKMİŞ", "Tarihsiz") else SARI if durum == "Bugün" else PatternFill()
        t.cell(t.max_row, 9).fill = SARI
        for c in t[t.max_row]:
            c.alignment = UST

    k = wb.create_sheet("Görüşme Kayıtları")
    _baslik(k, ["Görüşme No", "Tarih", "Kanal", "Temsilci", "Müşteri", "Önceki Görüşmeler", "Kategori", "Talep", "Yapılan İşlem", "Durum",
                "Aksiyonlar", "Verilen Sözler", "Duygu", "Risk", "Kontrol", "Orijinal Not"],
            (11, 15, 11, 12, 16, 14, 18, 40, 40, 11, 50, 40, 12, 30, 34, 60))
    for g in gorusmeler:
        s = sonuc.get(g["id"])
        if not s:
            k.append([g["id"], g["zaman"], g["kanal"], g["temsilci"], g["musteri"], ", ".join(g["onceki"]), "", "AI yanıt vermedi"] + [""] * 7 + [g["not"]])
            continue
        aks = "\n".join(f"• {a['is']} — {a['sorumlu']}" + (f" — {a['son_tarih']:%d.%m.%Y}" if a.get("son_tarih") else "") for a in s["aksiyonlar"])
        soz = "\n".join(f"• {x['soz']}" + (f" — {x['son_tarih']:%d.%m.%Y}" if x.get("son_tarih") else "") for x in s["taahhutler"])
        k.append([g["id"], g["zaman"], g["kanal"], g["temsilci"], g["musteri"], ", ".join(g["onceki"]), s["kategori"], s["talep"], s["yapilan_islem"],
                  DURUM[s["durum"]], aks, soz, DUYGU[s["duygu"]], s["risk"], "; ".join(kontroller.get(g["id"], [])), g["not"]])
        k.cell(k.max_row, 2).number_format = "DD.MM.YYYY HH:MM"
        for c in (7, 8, 9, 11, 12, 13, 14):
            k.cell(k.max_row, c).fill = AI_DOLGU
        if kontroller.get(g["id"]):
            k.cell(k.max_row, 15).fill = KIRMIZI
        if s["risk"]:
            k.cell(k.max_row, 14).fill = KIRMIZI
        for c in k[k.max_row]:
            c.alignment = UST
    k.auto_filter.ref = k.dimensions

    b = wb.create_sheet("Bilgi")
    for satir in [["Tarih hesabı", "Süre ifadeleri koddan tarihe çevrilir. 'N iş günü': hafta sonu ve sabit resmî tatiller (1 Ocak, 23 Nisan, 1 Mayıs, "
                                   "19 Mayıs, 15 Temmuz, 30 Ağustos, 29 Ekim) atlanır; dini bayramlar ve arife yarım günleri --tatiller ile verilmelidir. "
                                   "'Bu hafta': o haftanın son iş günü. 'Haftaya': sonraki haftanın son iş günü."],
                  ["Gizlilik", "Müşteri adları modele takma adla (K1, K2) gönderildi; telefon/e-posta/IBAN/TCKN maskelendi"],
                  ["Model", llm.kullanim_ozeti()], ["Oluşturulma", datetime.now().strftime("%d.%m.%Y %H:%M")]]:
        b.append(satir)
    b.column_dimensions["A"].width = 16
    b.column_dimensions["B"].width = 130
    for c in b["B"]:
        c.alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(girdi: Path, cikti: Path, tatil_yolu: Path | None = None, bugun: date | None = None, evet: bool = False) -> dict:
    gorusmeler = sorted(gorusmeleri_oku(girdi), key=lambda g: g["zaman"] or datetime.min)
    if any(g["zaman"] is None for g in gorusmeler):
        raise llm.LLMHatasi("Tarihi okunamayan görüşme var (GG.AA.YYYY SS:DD bekleniyor).")
    ekstra = set()
    if tatil_yolu:
        for x in tatil_yolu.read_text(encoding="utf-8").split():
            t = zaman(x)
            if t:
                ekstra.add(t.date())
    harita, gorulen = OrderedDict(), {}
    for g in gorusmeler:
        if g["musteri"] and g["musteri"] not in harita:
            harita[g["musteri"]] = f"K{len(harita) + 1}"
        g["onceki"] = gorulen.get(g["musteri"], []) if g["musteri"] else []
        if g["musteri"]:
            gorulen[g["musteri"]] = g["onceki"] + [g["id"]]
    tekrar = sum(1 for g in gorusmeler if g["onceki"])
    print(f"[OK] {len(gorusmeler)} görüşme · {len(harita)} müşteri · tekrar arayan: {tekrar}")
    llm.onay_al(f"{len(gorusmeler)} görüşme notu (müşteri adları takma adlı; telefon/e-posta/IBAN/TCKN maskeli) gönderilecek.", evet)
    sonuc = ozetle(gorusmeler, harita)
    kontroller = {g["id"]: kontrol_et(g, sonuc[g["id"]], ekstra) for g in gorusmeler if g["id"] in sonuc}
    rapor_yaz(cikti, gorusmeler, sonuc, kontroller, bugun or date.today())
    return {"gorusmeler": gorusmeler, "sonuc": sonuc, "kontroller": kontroller, "harita": harita}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Görüşme notlarını standart kayda ve tarihli takip listesine çevirir.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "gorusmeler.csv",
                   help="Görüşmeler (.xlsx/.csv): Tarih, Not; Görüşme No, Kanal, Temsilci, Müşteri isteğe bağlı")
    p.add_argument("--tatiller", type=Path, help="Ek tatil günleri (.txt, GG.AA.YYYY, boşlukla/satırla ayrılmış) — dini bayramlar")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "gorusme_kayitlari.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.cikti, a.tatiller, evet=a.evet)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    bekleyen = sum(1 for x in s["sonuc"].values() if x["durum"] != "cozuldu")
    print(f"[OK] {len(s['sonuc'])}/{len(s['gorusmeler'])} görüşme kayda çevrildi · bekleyen/eskalasyon: {bekleyen}")
    for gid, k in s["kontroller"].items():
        for x in k:
            print(f"[!] {gid}: {x}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
