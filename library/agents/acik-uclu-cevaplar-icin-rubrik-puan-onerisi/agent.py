"""
Açık Uçlu Cevaplar İçin Rubrik Puan Önerisi — Gerisi Bende AI Agent
Eğitim › Akademik › Ölçme Değerlendirme Uzmanı

Açık uçlu sınav cevaplarını rubriğe (dereceli puanlama anahtarı) göre puanlamak için öneri hazırlar; puanı kod toplar,
son kararı öğretmen verir.
1. Kod soruları (Soru; Soru Metni; Kazanım; Beklenen Cevap), rubriği (Soru; Ölçüt; Puan; Tanım) ve cevapları
   (Öğrenci No; Soru; Cevap) okur. Öğrenci adı / sınıf sütunları gönderilmez; telefon, e-posta, TCKN maskelenir. Boş cevap
   modele gönderilmez, kod 0 puan verir.
2. Model her ölçüt için rubrikteki bir puan düzeyini, cevaptan kelimesi kelimesine kanıtı ve gerekçeyi seçer; kısa geri
   bildirim yazar.
3. Kod denetler: her ölçüt bir kez mi; puan rubrikteki düzeylerden biri mi; en düşük düzeyin üstündeki puanın kanıtı
   cevapta aynen var mı; gerekçe boş mu; **aynı soruya aynı cevabı veren öğrencilere aynı puan verilmiş mi**. Sorun
   varsa yalnız o cevap modele geri gönderilir.
4. Kod puanları toplar, öğrenci × soru tablosunu ve kazanım bazında sınıf başarı yüzdesini hesaplar.

Kullanım:
    python agent.py                                                           # örnek 2 soru, 3 öğrenci
    python agent.py --sorular sorular.xlsx --rubrik rubrik.xlsx --cevaplar cevaplar.xlsx
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import llm

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def sayi(x) -> Decimal | None:
    s = str(x if x is not None else "").strip().replace(",", ".")
    try:
        return Decimal(s) if s else None
    except InvalidOperation:
        return None


def g(x: Decimal) -> int | float:
    return int(x) if x == x.to_integral_value() else float(x)


# ----------------------------------------------------------------------------
# Girdi
# ----------------------------------------------------------------------------

def tablo_oku(yol: Path) -> list[dict]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, read_only=True, data_only=True)
        s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=metin.split("\n")[0].count)))
    s = [r for r in s if any(c not in (None, "") for c in r)]
    bas = [katla(c) for c in s[0]] if s else []
    return [{k: ("" if v is None else str(v).strip()) for k, v in zip(bas, list(r) + [None] * (len(bas) - len(r)))} for r in s[1:]]


def soru_no(x) -> str:
    v = sayi(x)
    return str(g(v)) if v is not None else str(x).strip()


def oku(s_yol: Path, r_yol: Path, c_yol: Path) -> tuple[dict, list[dict]]:
    sorular = {}
    for r in tablo_oku(s_yol):
        if r.get("soru"):
            sorular[soru_no(r["soru"])] = {"no": soru_no(r["soru"]), "metin": r.get("soru metni", ""), "kazanim": r.get("kazanim", ""),
                                           "beklenen": r.get("beklenen cevap", ""), "olcutler": {}}
    for r in tablo_oku(r_yol):
        n, o, p = soru_no(r.get("soru", "")), r.get("olcut", ""), sayi(r.get("puan"))
        if n not in sorular:
            raise llm.LLMHatasi(f"{r_yol.name}: rubrikteki soru {n} soru listesinde yok.")
        if not o or p is None:
            continue
        sorular[n]["olcutler"].setdefault(o, {})[p] = r.get("tanim", "")
    eksik = [n for n, s in sorular.items() if not s["olcutler"]]
    if eksik:
        raise llm.LLMHatasi(f"Rubriği olmayan soru: {', '.join(eksik)}")
    cevaplar = []
    for r in tablo_oku(c_yol):
        o, n = r.get("ogrenci no") or r.get("ogrenci"), soru_no(r.get("soru", ""))
        if not o or not n:
            continue
        if n not in sorular:
            raise llm.LLMHatasi(f"{c_yol.name}: öğrenci {o} için soru {n} soru listesinde yok.")
        cevaplar.append({"id": f"{o}/{n}", "ogrenci": o, "soru": n, "cevap": r.get("cevap", ""), "cevap_m": llm.maskele(r.get("cevap", ""))})
    if not cevaplar:
        raise llm.LLMHatasi(f"{c_yol.name}: 'Öğrenci No; Soru; Cevap' sütunlarıyla en az bir cevap gerekli.")
    if len({c["id"] for c in cevaplar}) != len(cevaplar):
        raise llm.LLMHatasi(f"{c_yol.name}: aynı öğrenci ve soru için birden çok cevap var.")
    return sorular, cevaplar


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

SEMA = {"type": "object", "properties": {
    "olcutler": {"type": "array", "items": {"type": "object", "properties": {"olcut": S, "puan": S, "kanit": S, "gerekce": S},
                                            "required": ["olcut", "puan", "kanit", "gerekce"], "additionalProperties": False}},
    "geri_bildirim": S, "ogretmene_not": S},
    "required": ["olcutler", "geri_bildirim", "ogretmene_not"], "additionalProperties": False}


def toplam(c: dict) -> Decimal:
    return sum((sayi(x["puan"]) or Decimal(0) for x in c["olcutler"]), Decimal(0))


def denetle(cv: dict, soru: dict, c: dict) -> list[str]:
    on, kc, sorun = cv["id"], katla(cv["cevap_m"]), []
    adlar = [x["olcut"] for x in c["olcutler"]]
    eksik = [o for o in soru["olcutler"] if o not in adlar]
    if eksik:
        sorun.append(f"{on}: puanlanmayan ölçüt: {', '.join(eksik)}")
    fazla = sorted({o for o in adlar if o not in soru["olcutler"]} | {o for o in adlar if adlar.count(o) > 1})
    if fazla:
        sorun.append(f"{on}: rubrikte olmayan ya da tekrar eden ölçüt: {', '.join(fazla)}")
    for x in c["olcutler"]:
        duz = soru["olcutler"].get(x["olcut"])
        if duz is None:
            continue
        p = sayi(x["puan"])
        if p not in duz:
            sorun.append(f"{on}: {x['olcut']}: puan '{x['puan']}' rubrikte yok; şunlardan biri olmalı: {', '.join(str(g(d)) for d in sorted(duz))}")
            continue
        kk = katla(x["kanit"])
        if p > min(duz):
            if len(kk.split()) < 2 or kk not in kc:
                sorun.append(f"{on}: {x['olcut']}: {g(p)} puanın kanıtı cevapta aynen geçmiyor; öğrencinin yazdığını kelimesi kelimesine alın")
        elif kk and kk not in kc:
            sorun.append(f"{on}: {x['olcut']}: kanıt cevapta aynen geçmiyor")
        if not x["gerekce"].strip():
            sorun.append(f"{on}: {x['olcut']}: gerekçe boş")
    return sorun


def tutarlilik(cevaplar: list[dict], sonuc: dict) -> list[str]:
    """Aynı soruya aynı (katlanmış) cevabı veren öğrencilerin ölçüt puanları aynı olmalı."""
    ilk, sorun = {}, []
    for cv in cevaplar:
        if cv["id"] not in sonuc or not katla(cv["cevap"]):
            continue
        anahtar = (cv["soru"], katla(cv["cevap"]))
        p = {x["olcut"]: sayi(x["puan"]) for x in sonuc[cv["id"]]["olcutler"]}
        if anahtar in ilk:
            o, q = ilk[anahtar]
            if p != q:
                sorun.append(f"{cv['id']}: aynı cevap {o} için " + ", ".join(f"{k} {g(v)}" for k, v in q.items() if v is not None)
                             + " puan aldı; bu cevap " + ", ".join(f"{k} {g(v)}" for k, v in p.items() if v is not None) + ". Tutarlı puan verin")
        else:
            ilk[anahtar] = (cv["id"], p)
    return sorun


def rubrik_metni(s: dict) -> str:
    return "\n".join(f"{o} ({', '.join(f'{g(p)}: {t}' for p, t in sorted(d.items()))})" for o, d in s["olcutler"].items())


def uret(cevaplar: list[dict], sorular: dict, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    sonuc = dict(onceki or {})
    hedef = cevaplar if duzeltme is None else [c for c in cevaplar if any(s.startswith(c["id"] + ":") for s in duzeltme)]
    for cv in hedef:
        s = sorular[cv["soru"]]
        mesaj = (f"<soru no=\"{s['no']}\">{s['metin']}\nBeklenen cevap: {s['beklenen'] or '-'}</soru>\n<rubrik>\n{rubrik_metni(s)}\n</rubrik>\n"
                 f"<cevap>{cv['cevap_m']}</cevap>")
        d = [x.split(": ", 1)[1] for x in (duzeltme or []) if x.startswith(cv["id"] + ":")]
        if d:
            mesaj += ("\n<duzeltme>\nÖnceki cevap: " + json.dumps(sonuc.get(cv["id"]), ensure_ascii=False) + "\nKodun bulduğu sorunlar:\n"
                      + "\n".join(f"- {x}" for x in d) + "\n</duzeltme>")
        sonuc[cv["id"]] = llm.json_iste(sistem, mesaj, SEMA)
    return sonuc


def duzeltme_dongusu(uret_f, denetle_f, tur: int) -> tuple[dict, list[str], int]:
    taslak = uret_f(None, None)
    kalan = denetle_f(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret_f(kalan, taslak)
        kalan = denetle_f(taslak)
    return taslak, kalan, yapilan


# ----------------------------------------------------------------------------
# Hesap ve çıktı
# ----------------------------------------------------------------------------

def azami(s: dict) -> Decimal:
    return sum((max(d) for d in s["olcutler"].values()), Decimal(0))


def puanlar(cevaplar: list[dict], sorular: dict, sonuc: dict) -> dict:
    return {cv["id"]: (toplam(sonuc[cv["id"]]) if cv["id"] in sonuc else Decimal(0)) for cv in cevaplar}


def kazanim_ozeti(cevaplar: list[dict], sorular: dict, p: dict) -> list[list]:
    t = defaultdict(lambda: [Decimal(0), Decimal(0), 0])
    for cv in cevaplar:
        k = sorular[cv["soru"]]["kazanim"] or f"Soru {cv['soru']}"
        t[k][0] += p[cv["id"]]
        t[k][1] += azami(sorular[cv["soru"]])
        t[k][2] += 1
    return [[k, v[2], g(v[0]), g(v[1]), float(round(v[0] / v[1] * 100, 1)) if v[1] else None] for k, v in t.items()]


def tablo(ws, bas: list[str], satirlar: list[list], ai: set[int] = frozenset(), genis: dict | None = None) -> None:
    ws.append(bas)
    for h in ws[1]:
        h.fill, h.font, h.alignment = BASLIK_DOLGU, BASLIK_YAZI, UST
    for r in satirlar:
        ws.append(r)
        for j, v in enumerate(r, 1):
            ws.cell(ws.max_row, j).alignment = UST
            if j in ai and v not in (None, ""):
                ws.cell(ws.max_row, j).fill = AI_DOLGU
    for j in range(1, len(bas) + 1):
        ws.column_dimensions[get_column_letter(j)].width = (genis or {}).get(j, 14)
    ws.freeze_panes = "A2"


def rapor_yaz(cikti: Path, sorular: dict, cevaplar: list[dict], sonuc: dict, p: dict, kalan: list[str], tur: int) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Puan Önerileri"
    rr = []
    for cv in cevaplar:
        c = sonuc.get(cv["id"])
        sor = "; ".join(x.split(": ", 1)[1] for x in kalan if x.startswith(cv["id"] + ":"))
        if c is None:
            rr.append([cv["ogrenci"], cv["soru"], "(boş)", "", 0, g(azami(sorular[cv["soru"]])), "", "Boş cevap: kod 0 puan verdi", "", "", "Geçti"])
            continue
        det = "\n".join(f"{x['olcut']}: {x['puan']} — {x['gerekce']}" + (f" [“{x['kanit']}”]" if x["kanit"] else "") for x in c["olcutler"])
        rr.append([cv["ogrenci"], cv["soru"], cv["cevap"], det, g(p[cv["id"]]), g(azami(sorular[cv["soru"]])), "", c["geri_bildirim"], c["ogretmene_not"], "",
                   sor or "Geçti"])
    tablo(ws, ["Öğrenci No", "Soru", "Cevap", "Ölçüt Puanları (öneri)", "Önerilen Puan", "Azami", "Öğretmen Puanı", "Geri Bildirim", "Öğretmene Not",
               "Öğretmen Açıklaması", "Son Kontrol"], rr, ai={4, 8, 9}, genis={3: 60, 4: 70, 8: 45, 9: 30, 10: 30, 11: 30})
    for i, r in enumerate(rr, 2):
        ws.cell(i, 7).fill = SARI
        ws.cell(i, 11).fill = YESIL if r[10] == "Geçti" else KIRMIZI
    g2 = wb.create_sheet("Öğrenci × Soru")
    ogr = list(dict.fromkeys(cv["ogrenci"] for cv in cevaplar))
    sn = list(sorular)
    pm = {(cv["ogrenci"], cv["soru"]): p[cv["id"]] for cv in cevaplar}
    tablo(g2, ["Öğrenci No"] + [f"Soru {s} (/{g(azami(sorular[s]))})" for s in sn] + ["Toplam", "Azami"],
          [[o] + [g(pm[(o, s)]) if (o, s) in pm else None for s in sn]
           + [g(sum((pm.get((o, s), Decimal(0)) for s in sn), Decimal(0))), g(sum((azami(sorular[s]) for s in sn if (o, s) in pm), Decimal(0)))] for o in ogr])
    g3 = wb.create_sheet("Kazanım Özeti")
    tablo(g3, ["Kazanım", "Cevap Sayısı", "Alınan Puan", "Azami Puan", "Başarı %"], kazanim_ozeti(cevaplar, sorular, p), genis={1: 34})
    k = wb.create_sheet("Kontrol")
    k.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    k.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    k.cell(2, 1).fill = KIRMIZI if kalan else YESIL
    k.append(["Puanlar ÖNERİDİR. Not kararını öğretmen verir: 'Öğretmen Puanı' sütununu doldurun; öneriden ayrıldığınız yerleri açıklayın. "
              "Mor hücreler modelin; toplamlar, öğrenci × soru tablosu ve kazanım yüzdeleri kodun hesabıdır (öneri puanlarıyla). Öğrenci adı "
              "modele gönderilmez; öğrenci numarasını da gerekirse takma numarayla verin."])
    k.column_dimensions["A"].width = 140
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(s_yol: Path, r_yol: Path, c_yol: Path, cikti: Path, tur: int = 2, evet: bool = False) -> dict:
    sorular, cevaplar = oku(s_yol, r_yol, c_yol)
    dolu = [c for c in cevaplar if katla(c["cevap"])]
    print(f"[OK] {len(sorular)} soru · {len(cevaplar)} cevap · boş {len(cevaplar) - len(dolu)} (kod 0 puan verir)")
    sonuc, kalan, yapilan = {}, [], 0
    if dolu:
        llm.onay_al(f"{len(dolu)} cevap, soru metinleri ve rubrik gönderilecek (öğrenci adı ve numarası gönderilmez; telefon, e-posta, TCKN maskeli).", evet)
        sonuc, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(dolu, sorular, d, o),
                                                 lambda x: [s for c in dolu for s in denetle(c, sorular[c["soru"]], x[c["id"]])] + tutarlilik(dolu, x), tur)
    p = puanlar(cevaplar, sorular, sonuc)
    rapor_yaz(cikti, sorular, cevaplar, sonuc, p, kalan, yapilan)
    return {"sorular": sorular, "cevaplar": cevaplar, "sonuc": sonuc, "puan": p, "kalan": kalan, "tur": yapilan}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Açık uçlu cevaplar için rubriğe dayalı puan önerisi hazırlar; puanı ve tutarlılığı kod denetler.")
    p.add_argument("--sorular", type=Path, default=ORNEK / "sorular.csv", help="Soru; Soru Metni; Kazanım; Beklenen Cevap")
    p.add_argument("--rubrik", type=Path, default=ORNEK / "rubrik.csv", help="Soru; Ölçüt; Puan; Tanım (her puan düzeyi bir satır)")
    p.add_argument("--cevaplar", type=Path, default=ORNEK / "cevaplar.csv", help="Öğrenci No; Soru; Cevap")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "rubrik_puan_onerileri.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.sorular, a.rubrik, a.cevaplar, a.cikti, a.tur, a.evet)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    for cv in s["cevaplar"]:
        print(f"[OK] {cv['id']} · önerilen {g(s['puan'][cv['id']])} / {g(azami(s['sorular'][cv['soru']]))}" + ("" if cv["id"] in s["sonuc"] else " (boş cevap)"))
    for x in s["kalan"]:
        print(f"[!] Kalan: {x}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
