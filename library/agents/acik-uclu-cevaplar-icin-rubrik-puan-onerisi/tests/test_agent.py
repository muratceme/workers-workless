"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import json
import os
import re
import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
CEVAP = json.loads((Path(__file__).parent / "sahte_cevaplar.json").read_text(encoding="utf-8"))
MESAJLAR = []


def hazir():
    return agent.oku(ORNEK / "sorular.csv", ORNEK / "rubrik.csv", ORNEK / "cevaplar.csv")


SORU, CEVAPLAR = hazir()
METINDEN = {}
for _c in CEVAPLAR:
    METINDEN.setdefault((_c["soru"], _c["cevap"]), []).append(_c["id"])


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    """103/2'de ilk turda 101/2 ile birebir aynı cevaba farklı puan verir."""
    MESAJLAR.append(kullanici)
    soru = re.search(r'<soru no="([^"]+)"', kullanici)[1]
    cevap = re.search(r"<cevap>(.*?)</cevap>", kullanici, re.S)[1]
    ids = METINDEN[(soru, cevap)]
    if len(ids) > 1:                                                     # aynı cevap: sırayla 101/2, sonra 103/2
        sayac = sum(1 for m in MESAJLAR if f"<cevap>{cevap}</cevap>" in m and "<duzeltme>" not in m)
        if "<duzeltme>" in kullanici:
            return copy.deepcopy(CEVAP["103/2"])
        return copy.deepcopy(CEVAP["101/2" if sayac == 1 else "103/2_hatali"])
    return copy.deepcopy(CEVAP[ids[0]])


class KodTesti(unittest.TestCase):
    def test_okuma(self):
        s, c = hazir()
        self.assertEqual(list(s), ["1", "2"])
        self.assertEqual(s["1"]["olcutler"]["Etmenler"][D(2)], "Üç etmen de doğru (ışık, su, karbondioksit)")
        self.assertEqual((agent.azami(s["1"]), agent.azami(s["2"])), (D(4), D(3)))
        self.assertEqual([x["id"] for x in c], ["101/1", "101/2", "102/1", "102/2", "103/1", "103/2"])
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "c.csv"
            y.write_text("Öğrenci No;Soru;Cevap\n101;9;x\n", encoding="utf-8")
            with self.assertRaises(llm.LLMHatasi):
                agent.oku(ORNEK / "sorular.csv", ORNEK / "rubrik.csv", y)

    def test_denetim(self):
        s, c = hazir()
        b = {x["id"]: x for x in c}
        self.assertEqual([agent.denetle(b[i], s[b[i]["soru"]], CEVAP[i]) for i in ("101/1", "101/2", "102/1", "103/1", "103/2")], [[]] * 5)
        x = copy.deepcopy(CEVAP["102/1"])
        x["olcutler"][0].update(puan="1,5")
        x["olcutler"][1].update(puan="1", kanit="Su bitkiye besin taşır", gerekce="")
        x["olcutler"].append(copy.deepcopy(x["olcutler"][0]) | {"olcut": "Yazım"})
        r = agent.denetle(b["102/1"], s["1"], x)
        for p in ("puan '1,5' rubrikte yok; şunlardan biri olmalı: 0, 1, 2", "1 puanın kanıtı cevapta aynen geçmiyor", "gerekçe boş", "tekrar eden ölçüt: Yazım"):
            self.assertTrue(any(p in y for y in r), p)
        x = copy.deepcopy(CEVAP["101/2"])
        x["olcutler"].pop()
        self.assertEqual(agent.denetle(b["101/2"], s["2"], x), ["101/2: puanlanmayan ölçüt: Örnek"])

    def test_tutarlilik_ve_kazanim(self):
        s, c = hazir()
        d = [x for x in c if agent.katla(x["cevap"])]
        son = {i: CEVAP[i] for i in ("101/1", "101/2", "102/1", "103/1")} | {"103/2": CEVAP["103/2_hatali"]}
        self.assertEqual(agent.tutarlilik(d, son),
                         ["103/2: aynı cevap 101/2 için Kavram farkı 1, Örnek 0 puan aldı; bu cevap Kavram farkı 2, Örnek 1. Tutarlı puan verin"])
        son["103/2"] = CEVAP["103/2"]
        self.assertEqual(agent.tutarlilik(d, son), [])
        p = agent.puanlar(c, s, son)
        self.assertEqual((p["101/1"], p["102/1"], p["102/2"], p["103/2"]), (D(4), D(1), D(0), D(1)))
        self.assertEqual(agent.kazanim_ozeti(c, s, p), [["F.7.4 · Fotosentez", 3, 9, 12, 75.0], ["F.6.4 · Isı ve sıcaklık", 3, 2, 9, 22.2]])


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ornek_ve_duzeltme(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "r.xlsx"
            self.assertEqual(agent.main(["--evet", "--cikti", str(cikti)]), 0)
            self.assertEqual(len(MESAJLAR), 6)                          # 5 dolu cevap + 103/2 düzeltmesi; boş cevap gönderilmez
            for m in MESAJLAR:
                self.assertNotIn("101", m.split("<cevap>")[0])          # öğrenci no gönderilmez
            wb = load_workbook(cikti)
            r = [[c.value for c in x] for x in wb["Puan Önerileri"].iter_rows(min_row=2)]
            self.assertEqual([(x[0], x[1], x[4], x[10]) for x in r], [("101", "1", 4, "Geçti"), ("101", "2", 1, "Geçti"), ("102", "1", 1, "Geçti"),
                                                                       ("102", "2", 0, "Geçti"), ("103", "1", 4, "Geçti"), ("103", "2", 1, "Geçti")])
            m = [[c.value for c in x] for x in wb["Öğrenci × Soru"].iter_rows()]
            self.assertEqual(m[1:], [["101", 4, 1, 5, 7], ["102", 1, 0, 1, 7], ["103", 4, 1, 5, 7]])

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(agent.main(["--rubrik", str(Path(tmp) / "yok.csv"), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
