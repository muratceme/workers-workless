"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

DOSYA = KLASOR / "ornek_veri" / "ornek_metal"
BUGUN = date(2026, 10, 8)
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return {"firma_kunyesi": "[FİRMA] 2011'de kurulmuş, sermayesi 25.000.000 TL.",
            "ortaklik_ve_yonetim": "[KİŞİ-1] %60, [KİŞİ-2] %25, [ŞİRKET-1] %10 paya sahip.",
            "faaliyet_ve_piyasa": "Sac metal işleme; tedarikçiye göre ödemeler sarkıyor.",
            "banka_iliskileri": "Örnek Bank B istihbaratında 7.200.000 TL risk bildirildi.",
            "olumsuz_bilgiler": "Açık bir icra takibi var.",
            "celiskiler": [{"aciklama": "Bank B riski istihbaratta ve KKB'de farklı.", "kaynak": "banka istihbaratı, KKB"}],
            "guclu_yonler": [{"aciklama": "Müşteri kaliteden memnun.", "kaynak": "piyasa görüşmeleri"}],
            "risk_isaretleri": [{"aciklama": "Tedarikçi bakiyesi 4.800.000 TL.", "onem": "orta", "kaynak": "piyasa"}],
            "teyit_edilecekler": ["Ortaklık payının eksik %5'i."],
            "genel_degerlendirme": "[FİRMA] büyüyen ancak likiditesi sıkışan bir görünümde."}


class KontrolTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.g, cls.m, cls.a = agent.dosyalari_oku(DOSYA)
        cls.k = agent.kontrol_et(cls.g, cls.m, BUGUN)
        cls.metin = "\n".join(f"{o} {a}" for o, a, _ in cls.k["bulgular"])

    def test_sicil(self):
        s = self.k["sicil"]
        self.assertEqual((s["unvan"], s["vkn"], s["kurulus"], s["sermaye"]),
                         ("Örnek Metal Sanayi ve Ticaret A.Ş.", "1234567890", date(2011, 3, 14), D(25000000)))
        self.assertEqual([(o["ad"], o["pay"], o["tuzel"]) for o in s["ortaklar"]],
                         [("Hasan Kurgucu", D(60), False), ("Leyla Kurgucu", D(25), False), ("Kurgu Holding A.Ş.", D(10), True)])
        self.assertEqual([d["tur"] for d in s["degisiklikler"]], ["Sermaye", "Yönetim", "Adres", "Unvan"])

    def test_bulgular(self):
        self.assertIn("orta Ortaklık payları toplamı %95", self.metin)
        self.assertIn("orta Son 12 ayda 3 sicil değişikliği", self.metin)
        self.assertIn("orta Örnek Bank D KKB'de görünüyor", self.metin)
        self.assertIn("istihbarattaki nakdi risk 7.200.000 TL, KKB'de 7.900.000 TL", self.metin)
        self.assertIn("yüksek İcra takibi 10.09.2026", self.metin)
        self.assertIn("orta Karşılıksız çek 14.04.2024", self.metin)            # 3 yıl içinde, ödenmiş
        self.assertIn("bilgi Protestolu senet 05.01.2021", self.metin)
        self.assertIn("'sark': Son 6 ayda ödemeler", self.metin)
        self.assertNotIn("anlaşmazlık", self.metin)                              # "duyulmadı" ile olumsuzlanmış
        self.assertNotIn("bulunmamaktadır", self.metin)
        self.assertEqual(self.k["risk"]["nr"], D(26450000))

    def test_takma_adlar(self):
        h = agent.takma_adlar(self.k["sicil"], ["Ali Görüşülen"])
        self.assertEqual(h["Kurgu Holding A.Ş."], "[ŞİRKET-1]")
        self.assertEqual(h["Okan Denemeci"], "[KİŞİ-3]")
        self.assertEqual(h["Ali Görüşülen"], "[KİŞİ-4]")
        self.assertEqual(h["Kurgucu"], "[KİŞİ-1-SOYAD]")


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_uctan_uca(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "i.xlsx"
            s = agent.calistir(DOSYA, cikti, ["Selim Taslakoğlu"], BUGUN, evet=True)
            m = MESAJLAR[0]
            for gizli in ("Örnek Metal", "1234567890", "Hasan", "Kurgucu", "Okan Denemeci", "Kurgu Holding", "Taslakoğlu", "0532 000 44 55"):
                self.assertNotIn(gizli, m)
            for var in ("[FİRMA]", "[KİŞİ-1]", "[ŞİRKET-1]", "<kod_kontrolleri>", "Örnek Bank D"):
                self.assertIn(var, m)
            self.assertEqual(s["yanit"]["dogrulanamayan_sayilar"], ["4800000"])     # kaynakta 4.500.000
            md = s["md"].read_text(encoding="utf-8")
            self.assertIn("Hasan Kurgucu %60, Leyla Kurgucu %25, Kurgu Holding A.Ş. %10", md)
            self.assertIn("Örnek Metal Sanayi ve Ticaret A.Ş. 2011", md)
            self.assertNotIn("[KİŞİ", md)
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Özet", "Kontroller", "Ortaklık ve Sicil", "Banka İstihbaratı", "KKB Risk", "Olumsuz Kayıtlar"])
            self.assertEqual(wb["KKB Risk"]["G4"].value, "Hayır")                  # Örnek Bank D

    def test_klasor_yok(self):
        with self.assertRaises(llm.LLMHatasi):
            agent.calistir(Path("olmayan_klasor"), Path("x.xlsx"), evet=True)


if __name__ == "__main__":
    unittest.main()
