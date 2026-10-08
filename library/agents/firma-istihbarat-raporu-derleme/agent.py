"""
Firma İstihbarat Raporu Derleme — Workers / Workless AI Agent
Bankacılık › Şube Bankacılığı › Ticari Portföy Yöneticisi

1. Klasördeki dosyalar okunur: ticaret sicili özeti (.txt/.pdf/.docx), banka istihbaratı, piyasa görüşmeleri ve
   diğer notlar; KKB/memzuç risk tablosu ve olumsuz kayıtlar (karşılıksız çek, protesto, icra) tabloları.
2. Kod kontrol eder: VKN, firma yaşı, ortaklık paylarının toplamı, tüzel kişi ortaklar, son 12 aydaki sicil
   değişiklikleri (sermaye, yönetim, adres, unvan), KKB'de görünüp istihbaratı alınmamış bankalar, istihbarattaki
   limit/risk ile KKB'nin farkı, banka bazında limit doluluğu, olumsuz kayıtların yaşı ve durumu, metinlerdeki
   olumsuz ifadeler ("bulunmamaktadır", "duyulmadı" gibi olumsuzlananlar sayılmaz).
3. Firma unvanı [FİRMA], VKN [VKN], sicildeki ortak ve yöneticiler [KİŞİ-n] / [ŞİRKET-n] olarak; telefon, e-posta,
   TCKN ve IBAN maskelenir. Model kaynaklara dayanarak istihbarat raporunu yazar; kaynaklar arası çelişkileri ve
   teyit edilecek noktaları listeler. Modelin yazdığı, girdilerde olmayan sayılar işaretlenir.
4. Çıktı: Markdown rapor (takma adlar geri açılmış) + Excel (Özet, Kontroller + 'İnceleme', Ortaklık ve Sicil,
   Banka İstihbaratı, KKB Risk, Olumsuz Kayıtlar).

Kullanım:
    python agent.py                                          # örnek: kurgusal metal işleme firması
    python agent.py --girdi ./istihbarat_klasoru --tarih 08.10.2026 --gizle "Görüşülen Kişi"
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
SIFIR = Decimal(0)
OLUMSUZ = [("karşılıksız", "yüksek"), ("protest", "yüksek"), ("haciz", "yüksek"), ("icra", "yüksek"), ("iflas", "yüksek"),
           ("konkordato", "yüksek"), ("yasal takip", "yüksek"), ("takibe", "yüksek"), ("gecikme", "orta"), ("sark", "orta"),
           ("yapılandır", "orta"), ("anlaşmazlık", "orta"), ("dava", "orta"), ("olumsuz", "orta"), ("geç öde", "orta")]
OLUMSUZLAMA = re.compile(r"\b(yok|yoktur|bulunmamaktadır|bulunmamakta|rastlanmamıştır|görülmemiştir|duyulmadı|olmamıştır|yaşanmamıştır)\b")
TUZEL = re.compile(r"\b(A\.?Ş\.?|Ltd\.?|Şti\.?|Holding|Anonim|Limited|Kooperatif|Vakf|Derneği)\b", re.I)


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def katla(s) -> str:
    s = kucuk(s)
    for a, b in zip("ıöüşğçâîû", "iousgcaiu"):
        s = s.replace(a, b)
    return re.sub(r"[^a-z0-9%]+", " ", s).strip()


def para(x) -> Decimal | None:
    if x in (None, "", "—", "-"):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("₺", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tl(x) -> str:
    return "—" if x is None else f"{x:,.0f}".replace(",", ".")


def yuzde(x) -> str:
    return "—" if x is None else f"%{x * 100:.1f}".replace(".", ",")


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(str(x or "").strip(), f).date()
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = belge.txt_oku(yol)
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def vkn_kontrol(no: str) -> bool:
    if not re.fullmatch(r"\d{10}", no or ""):
        return False
    s = 0
    for i, n in enumerate(reversed(no[:9]), 1):
        c1 = (int(n) + i) % 10
        if c1:
            s += (c1 * 2 ** i) % 9 or 9
    return (10 - s % 10) % 10 == int(no[9])


# ----------------------------------------------------------------------------
# Girdiler
# ----------------------------------------------------------------------------

@dataclass
class Dosya:
    ad: str
    tur: str
    metin: str


def kkb_oku(yol: Path) -> list[dict]:
    s = tablo_oku(yol)
    bi = next((i for i, r in enumerate(s[:10]) if any(katla(c) in ("banka", "kurum") for c in r)), 0)
    b = [katla(c) for c in s[bi]]
    i = lambda *adlar: next((b.index(a) for a in adlar if a in b), None)  # noqa: E731
    ib, inl, inr, igl, igr = i("banka", "kurum"), i("nakdi limit"), i("nakdi risk"), i("gayrinakdi limit", "gayri nakdi limit"), i("gayrinakdi risk", "gayri nakdi risk")
    sonuc = []
    for r in s[bi + 1:]:
        if not r[ib] or katla(r[ib]).startswith(("toplam", "genel")):
            continue
        al = lambda j: (para(r[j]) or SIFIR) if j is not None and j < len(r) else SIFIR  # noqa: E731
        sonuc.append({"banka": str(r[ib]).strip(), "nl": al(inl), "nr": al(inr), "gl": al(igl), "gr": al(igr)})
    return sonuc


def olumsuz_oku(yol: Path) -> list[dict]:
    s = tablo_oku(yol)
    b = [katla(c) for c in s[0]]
    i = lambda *adlar: next((b.index(a) for a in adlar if a in b), None)  # noqa: E731
    it, ita, itu, idu, ia = i("tur", "kayit turu"), i("tarih"), i("tutar"), i("durum"), i("aciklama")
    al = lambda r, j: r[j] if j is not None and j < len(r) else None  # noqa: E731
    return [{"tur": str(al(r, it) or "").strip(), "tarih": tarih(al(r, ita)), "tutar": para(al(r, itu)), "durum": str(al(r, idu) or "").strip(),
             "aciklama": str(al(r, ia) or "").strip()} for r in s[1:] if al(r, it)]


def dosyalari_oku(klasor: Path) -> tuple[dict, list[Dosya], list[str]]:
    girdi = {"sicil": None, "kkb": None, "olumsuz": None}
    metinler, atlanan = [], []
    for yol in sorted(klasor.iterdir()):
        ad = katla(yol.stem)
        if yol.suffix.lower() in {".csv", ".xlsx", ".xlsm"}:
            ilk = katla(" ".join(str(c) for r in tablo_oku(yol)[:2] for c in r))
            if "kkb" in ad or "memzuc" in ad or "nakdi" in ilk:
                girdi["kkb"], tur = kkb_oku(yol), "KKB / memzuç risk"
            elif "olumsuz" in ad or "cek" in ad or "protesto" in ad or ("tur" in ilk and "durum" in ilk):
                girdi["olumsuz"], tur = olumsuz_oku(yol), "Olumsuz kayıtlar"
            else:
                atlanan.append(f"{yol.name} (tanınmayan tablo)")
                continue
            metinler.append(Dosya(yol.name, tur, ""))
            continue
        if yol.suffix.lower() not in belge.DESTEKLENEN:
            atlanan.append(yol.name)
            continue
        try:
            metin = belge.metin_oku(yol)
        except belge.BelgeHatasi as h:
            atlanan.append(f"{yol.name} ({h})")
            continue
        bas = katla(metin[:300])
        if "sicil" in ad or "ticaret sicil" in bas or "ttsg" in bas:
            tur = "Ticaret sicili"
            girdi["sicil"] = metin
        elif "banka" in ad or "banka istihbarat" in bas:
            tur = "Banka istihbaratı"
        elif "piyasa" in ad or "piyasa" in bas:
            tur = "Piyasa görüşmeleri"
        else:
            tur = "Diğer not"
        metinler.append(Dosya(yol.name, tur, metin))
    return girdi, metinler, atlanan


def alan(metin: str, *etiketler: str) -> str | None:
    for et in etiketler:
        m = re.search(r"^\s*" + re.escape(et) + r"\s*:\s*(.+)$", metin or "", re.I | re.M)
        if m:
            return m.group(1).strip()
    return None


def bolum(metin: str, baslik: str) -> list[str]:
    """'Ortaklar:' gibi bir başlığın altındaki '- ...' satırları."""
    m = re.search(r"^\s*" + baslik + r"[^\n]*:\s*\n((?:\s*[-•*].*\n?)+)", metin or "", re.I | re.M)
    return [x.strip(" -•*\t") for x in m.group(1).splitlines() if x.strip()] if m else []


def sicil_coz(metin: str | None) -> dict:
    s = {"unvan": alan(metin, "Unvan", "Ticaret Unvanı"), "vkn": re.sub(r"\D", "", alan(metin, "VKN", "Vergi Kimlik No", "Vergi No") or ""),
         "kurulus": tarih(alan(metin, "Kuruluş Tarihi", "Tescil Tarihi")), "sermaye": para(re.sub(r"[^\d.,]", "", alan(metin, "Sermaye") or "")),
         "faaliyet": alan(metin, "Faaliyet Konusu", "Faaliyet"), "ortaklar": [], "yonetim": [], "degisiklikler": []}
    for x in bolum(metin, "Ortaklar"):
        m = re.match(r"(.+?)\s*[:\-–]\s*%?\s*([\d.,]+)\s*%?$", x)
        if m:
            s["ortaklar"].append({"ad": m.group(1).strip(), "pay": para(m.group(2)), "tuzel": bool(TUZEL.search(m.group(1)))})
    for x in bolum(metin, "Yönetim Kurulu") or bolum(metin, "Yönetim") or bolum(metin, "Müdürler"):
        m = re.match(r"(.+?)(?:\s*\((.+)\))?$", x)
        s["yonetim"].append({"ad": m.group(1).strip(), "gorev": (m.group(2) or "").strip()})
    for x in bolum(metin, "Değişiklikler"):
        m = re.match(r"(\d{2}\.\d{2}\.\d{4})\s*[:\-–]\s*(.+)", x)
        if m:
            k = katla(m.group(2))
            tur = ("Sermaye" if "sermaye" in k else "Yönetim" if ("yonetim" in k or "mudur" in k or "uyelig" in k) else "Adres" if "adres" in k
                   else "Unvan" if "unvan" in k else "Ortaklık" if ("hisse" in k or "pay devr" in k or "ortak" in k) else "Diğer")
            s["degisiklikler"].append({"tarih": tarih(m.group(1)), "tur": tur, "aciklama": m.group(2).strip()})
    return s


def banka_istihbarati(metinler: list[Dosya]) -> list[dict]:
    sonuc = []
    for d in metinler:
        if d.tur != "Banka istihbaratı":
            continue
        for satir in d.metin.splitlines():
            m = re.match(r"\s*([^:]{3,60}?(?:bank|banka|bankası)[^:]{0,20}):\s*(.+)", satir, re.I)
            if not m:
                continue
            g = m.group(2)
            nl = re.search(r"nakdi limit[^\d]{0,10}([\d.,]+)", g, re.I)
            nr = re.search(r"(?:nakdi )?risk[^\d]{0,10}([\d.,]+)", g, re.I)
            olumsuz = [k for k, _ in OLUMSUZ if k in kucuk(g) and not OLUMSUZLAMA.search(kucuk(g))]
            sonuc.append({"banka": m.group(1).strip(), "metin": g.strip(), "nl": para(nl.group(1).rstrip(".,")) if nl else None,
                          "nr": para(nr.group(1).rstrip(".,")) if nr else None, "olumsuz": olumsuz})
    return sonuc


def kontrol_et(girdi: dict, metinler: list[Dosya], bugun: date) -> dict:
    bulgular = []

    def b(onem, aciklama, kaynak=""):
        bulgular.append((onem, aciklama, kaynak))

    s = sicil_coz(girdi["sicil"])
    if not girdi["sicil"]:
        b("orta", "Ticaret sicili özeti yok: ortaklık, yönetim ve sermaye bilgisi doğrulanamadı")
    else:
        if s["vkn"] and not vkn_kontrol(s["vkn"]):
            b("yüksek", f"VKN kontrol hanesi tutmuyor: {s['vkn']}", "ticaret sicili")
        if s["kurulus"]:
            yas = (bugun - s["kurulus"]).days / 365.25
            s["yas"] = yas
            if yas < 3:
                b("orta", f"Firma {yas:.1f} yıllık (3 yıldan genç)".replace(".", ","), "ticaret sicili")
        toplam = sum((o["pay"] or SIFIR) for o in s["ortaklar"])
        if s["ortaklar"] and toplam != 100:
            b("orta", f"Ortaklık payları toplamı %{toplam:g} (100 olmalı): sicil özeti eksik veya güncel değil olabilir", "ticaret sicili")
        for o in s["ortaklar"]:
            if o["tuzel"]:
                b("bilgi", f"Tüzel kişi ortak: {o['ad']} (%{o['pay']:g}) — grup yapısı ve grup riski incelenmeli", "ticaret sicili")
        son12 = [d for d in s["degisiklikler"] if d["tarih"] and (bugun - d["tarih"]).days <= 365]
        if son12:
            b("orta" if len(son12) >= 3 else "bilgi", f"Son 12 ayda {len(son12)} sicil değişikliği: " +
              "; ".join(f"{d['tarih']:%d.%m.%Y} {d['tur'].lower()}" for d in son12), "ticaret sicili")
    kkb = girdi["kkb"] or []
    risk = {a: sum((x[a] for x in kkb), SIFIR) for a in ("nl", "nr", "gl", "gr")}
    istihbarat = banka_istihbarati(metinler)
    if kkb:
        if risk["nl"] and risk["nr"] / risk["nl"] >= Decimal("0.9"):
            b("orta", f"Sektörde nakdi limit doluluğu {yuzde(risk['nr'] / risk['nl'])} ({tl(risk['nr'])} / {tl(risk['nl'])} TL)", "KKB")
        for x in kkb:
            if x["nl"] and x["nr"] / x["nl"] >= Decimal("0.95"):
                b("bilgi", f"{x['banka']}: nakdi limit doluluğu {yuzde(x['nr'] / x['nl'])}", "KKB")
            ist = next((i for i in istihbarat if katla(i["banka"]) == katla(x["banka"])), None)
            if ist is None:
                b("orta", f"{x['banka']} KKB'de görünüyor ({tl(x['nr'])} TL nakdi risk) ama banka istihbaratı alınmamış", "KKB")
            elif ist["nr"] is not None and ist["nr"] != x["nr"]:
                b("bilgi", f"{x['banka']}: istihbarattaki nakdi risk {tl(ist['nr'])} TL, KKB'de {tl(x['nr'])} TL (tarih farkı olabilir)", "KKB / istihbarat")
    else:
        b("orta", "KKB / memzuç risk tablosu yok")
    for i in istihbarat:
        if kkb and i["nr"] and not any(katla(x["banka"]) == katla(i["banka"]) for x in kkb):
            b("bilgi", f"{i['banka']} istihbaratta risk bildiriyor ama KKB tablosunda yok", "istihbarat / KKB")
    ol = girdi["olumsuz"] or []
    for x in ol:
        acik = "acik" in katla(x["durum"]) or "devam" in katla(x["durum"])
        yas_gun = (bugun - x["tarih"]).days if x["tarih"] else None
        onem = "yüksek" if acik or (yas_gun is not None and yas_gun <= 365) else "orta" if yas_gun is not None and yas_gun <= 3 * 365 else "bilgi"
        b(onem, f"{x['tur']} {x['tarih']:%d.%m.%Y} · {tl(x['tutar'])} TL · {x['durum'] or 'durum yok'}" + (f" · {x['aciklama']}" if x["aciklama"] else ""),
          "olumsuz kayıtlar")
    for d in metinler:
        if d.tur not in ("Banka istihbaratı", "Piyasa görüşmeleri", "Diğer not"):
            continue
        for cumle in re.split(r"(?<=[.!?])\s+|\n", d.metin):
            ck = kucuk(cumle)
            for kelime, onem in OLUMSUZ:
                if kelime in ck and not OLUMSUZLAMA.search(ck):
                    b(onem, f"'{kelime}': {cumle.strip()[:170]}", d.ad)
                    break
    return {"bulgular": bulgular, "sicil": s, "kkb": kkb, "risk": risk, "istihbarat": istihbarat, "olumsuz": ol}


# ----------------------------------------------------------------------------
# Maskeleme
# ----------------------------------------------------------------------------

def takma_adlar(s: dict, terimler: list[str]) -> dict[str, str]:
    harita = {}
    if s.get("unvan"):
        harita[s["unvan"]] = "[FİRMA]"
        kok = re.split(r"\s+(?:San\.|Sanayi|Tic\.|Ticaret|A\.Ş\.|Ltd\.)", s["unvan"])[0].strip()
        if len(kok) > 4 and kok != s["unvan"]:
            harita[kok] = "[FİRMA]"
    if s.get("vkn"):
        harita[s["vkn"]] = "[VKN]"
    kisiler, sirketler = [], []
    for ad in [o["ad"] for o in s.get("ortaklar", [])] + [y["ad"] for y in s.get("yonetim", [])] + list(terimler):
        ad = ad.strip()
        if not ad or ad in harita:
            continue
        (sirketler if TUZEL.search(ad) else kisiler).append(ad)
    for i, ad in enumerate(dict.fromkeys(sirketler), 1):
        harita[ad] = f"[ŞİRKET-{i}]"
    for i, ad in enumerate(dict.fromkeys(kisiler), 1):
        harita[ad] = f"[KİŞİ-{i}]"
        parcalar = ad.split()
        if len(parcalar) > 1 and len(parcalar[-1]) > 3:
            harita.setdefault(parcalar[-1], f"[KİŞİ-{i}-SOYAD]")
    return harita


def maskele(metin: str, harita: dict[str, str]) -> str:
    for gercek, takma in sorted(harita.items(), key=lambda x: -len(x[0])):
        metin = re.sub(re.escape(gercek), takma, metin, flags=re.I)
    return llm.maskele(metin)


def geri_ac(metin: str, harita: dict[str, str]) -> str:
    for gercek, takma in sorted(harita.items(), key=lambda x: -len(x[1])):
        if takma != "[FİRMA]":
            metin = metin.replace(takma, gercek)
    unvan = next((g for g, t in harita.items() if t == "[FİRMA]"), None)
    return metin.replace("[FİRMA]", unvan) if unvan else metin


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

KAYNAKLI = {"type": "object", "properties": {"aciklama": {"type": "string"}, "kaynak": {"type": "string"}},
            "required": ["aciklama", "kaynak"], "additionalProperties": False}
SEMA = {
    "type": "object",
    "properties": {
        "firma_kunyesi": {"type": "string"},
        "ortaklik_ve_yonetim": {"type": "string"},
        "faaliyet_ve_piyasa": {"type": "string"},
        "banka_iliskileri": {"type": "string"},
        "olumsuz_bilgiler": {"type": "string"},
        "celiskiler": {"type": "array", "items": KAYNAKLI},
        "guclu_yonler": {"type": "array", "items": KAYNAKLI},
        "risk_isaretleri": {"type": "array", "items": {
            "type": "object", "properties": {"aciklama": {"type": "string"}, "onem": {"type": "string", "enum": ["yüksek", "orta", "düşük"]},
                                             "kaynak": {"type": "string"}},
            "required": ["aciklama", "onem", "kaynak"], "additionalProperties": False}},
        "teyit_edilecekler": {"type": "array", "items": {"type": "string"}},
        "genel_degerlendirme": {"type": "string"},
    },
    "required": ["firma_kunyesi", "ortaklik_ve_yonetim", "faaliyet_ve_piyasa", "banka_iliskileri", "olumsuz_bilgiler", "celiskiler",
                 "guclu_yonler", "risk_isaretleri", "teyit_edilecekler", "genel_degerlendirme"],
    "additionalProperties": False,
}
METIN_ALANLARI = ("firma_kunyesi", "ortaklik_ve_yonetim", "faaliyet_ve_piyasa", "banka_iliskileri", "olumsuz_bilgiler", "genel_degerlendirme")


def kod_ozeti(k: dict) -> list[str]:
    s, satir = k["sicil"], []
    if s["unvan"]:
        satir.append(f"- Unvan: {s['unvan']}; VKN: {s['vkn'] or '—'}; kuruluş: {s['kurulus']:%d.%m.%Y}" if s["kurulus"] else f"- Unvan: {s['unvan']}")
    if s.get("yas") is not None:
        satir.append(f"- Firma yaşı: {s['yas']:.1f} yıl".replace(".", ","))
    if s["sermaye"]:
        satir.append(f"- Sermaye: {tl(s['sermaye'])} TL")
    for o in s["ortaklar"]:
        satir.append(f"- Ortak: {o['ad']} %{o['pay']:g}" + (" (tüzel kişi)" if o["tuzel"] else ""))
    for y in s["yonetim"]:
        satir.append(f"- Yönetim: {y['ad']}" + (f" ({y['gorev']})" if y["gorev"] else ""))
    for d in s["degisiklikler"]:
        satir.append(f"- Sicil değişikliği {d['tarih']:%d.%m.%Y} ({d['tur']}): {d['aciklama']}")
    if k["kkb"]:
        r = k["risk"]
        satir.append(f"- KKB toplam: nakdi limit {tl(r['nl'])} / risk {tl(r['nr'])} TL, gayrinakdi limit {tl(r['gl'])} / risk {tl(r['gr'])} TL, "
                     f"{len(k['kkb'])} banka")
        satir += [f"- KKB {x['banka']}: nakdi {tl(x['nr'])} / {tl(x['nl'])} TL, gayrinakdi {tl(x['gr'])} / {tl(x['gl'])} TL" for x in k["kkb"]]
    for x in k["olumsuz"]:
        satir.append(f"- Olumsuz kayıt: {x['tur']} {x['tarih']:%d.%m.%Y}, {tl(x['tutar'])} TL, {x['durum']}" if x["tarih"] else f"- Olumsuz kayıt: {x['tur']}")
    return satir


def sayilar(metin: str) -> set[str]:
    """3+ haneli sayıların yalnız rakamları ('15.000.000' → '15000000')."""
    return {re.sub(r"\D", "", m) for m in re.findall(r"\d[\d.,]*\d", metin) if len(re.sub(r"\D", "", m)) >= 3}


def hazirla(metinler: list[Dosya], k: dict, harita: dict) -> tuple[str, str]:
    kk = [f"- [{o}] {maskele(a, harita)} (kaynak: {kay or '-'})" for o, a, kay in k["bulgular"]]
    mesaj = "\n".join([
        "<belgeler>",
        *[f'<belge ad="{d.ad}" tur="{d.tur}">\n{maskele(d.metin, harita)}\n</belge>' for d in metinler if d.metin],
        "</belgeler>",
        "<kod_ozeti>", *[maskele(x, harita) for x in kod_ozeti(k)], "</kod_ozeti>",
        "<kod_kontrolleri>", *(kk or ["- Kontrol bulgusu yok"]), "</kod_kontrolleri>",
    ])
    return (BURASI / "prompt.md").read_text(encoding="utf-8"), mesaj


def model_yaz(metinler, k, harita) -> dict:
    sistem, mesaj = hazirla(metinler, k, harita)
    yanit = llm.json_iste(sistem, mesaj, SEMA)
    bilinen = sayilar(mesaj) | {x for x in re.findall(r"\d+", mesaj) if len(x) >= 3}   # tarihlerin yıl kısmı ("14.03.2011" → 2011)
    metin = " ".join([yanit[a] for a in METIN_ALANLARI] + yanit["teyit_edilecekler"] +
                     [x["aciklama"] for a in ("celiskiler", "guclu_yonler", "risk_isaretleri") for x in yanit[a]])
    yanit["dogrulanamayan_sayilar"] = sorted(s for s in sayilar(metin) if s not in bilinen)
    return yanit


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
ONEM = {"yüksek": "FDE2E1", "orta": "FFF4CE", "bilgi": "E8F0FE", "düşük": "E8F0FE"}
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def rapor_yaz(cikti: Path, metinler, atlanan, k: dict, y: dict, harita: dict, bugun: date) -> Path:
    g = lambda s: geri_ac(s, harita)  # noqa: E731
    s = k["sicil"]
    wb = Workbook()
    o = wb.active
    o.title = "Özet"
    o.column_dimensions["A"].width, o.column_dimensions["B"].width = 30, 100
    for a, d in [("Firma", s["unvan"] or "—"), ("VKN", s["vkn"] or "—"), ("Rapor tarihi", f"{bugun:%d.%m.%Y}"),
                 ("Kaynaklar", "; ".join(f"{d.ad} ({d.tur})" for d in metinler)), ("Okunamayan", "; ".join(atlanan) or "—"),
                 ("Künye (model)", g(y["firma_kunyesi"])), ("Ortaklık ve yönetim (model)", g(y["ortaklik_ve_yonetim"])),
                 ("Faaliyet ve piyasa (model)", g(y["faaliyet_ve_piyasa"])), ("Banka ilişkileri (model)", g(y["banka_iliskileri"])),
                 ("Olumsuz bilgiler (model)", g(y["olumsuz_bilgiler"])), ("Genel değerlendirme (taslak)", g(y["genel_degerlendirme"])),
                 ("Doğrulanamayan sayılar", ", ".join(y["dogrulanamayan_sayilar"]) or "—"), ("Model", llm.kullanim_ozeti())]:
        o.append([a, d])
        o.cell(o.max_row, 1).font = Font(bold=True)
        if "(model)" in a or "taslak" in a:
            o.cell(o.max_row, 2).fill = PatternFill("solid", fgColor="EDE7FF")
        if a == "Doğrulanamayan sayılar" and y["dogrulanamayan_sayilar"]:
            o.cell(o.max_row, 2).fill = PatternFill("solid", fgColor=ONEM["yüksek"])
        for h in o[o.max_row]:
            h.alignment = UST

    kt = wb.create_sheet("Kontroller")
    _baslik(kt, ["Kaynak", "Önem", "Bulgu", "Dayanak", "İnceleme"], (8, 9, 90, 26, 26))
    for on, a, kay in k["bulgular"]:
        kt.append(["Kod", on, a, kay, ""])
        kt.cell(kt.max_row, 2).fill = PatternFill("solid", fgColor=ONEM[on])
    for x in y["risk_isaretleri"]:
        kt.append(["Model", x["onem"], g(x["aciklama"]), x["kaynak"], ""])
        kt.cell(kt.max_row, 2).fill = PatternFill("solid", fgColor=ONEM[x["onem"]])
    for x in y["celiskiler"]:
        kt.append(["Model", "çelişki", g(x["aciklama"]), x["kaynak"], ""])
    for t in y["teyit_edilecekler"]:
        kt.append(["Model", "teyit", g(t), "", ""])
    for r in kt.iter_rows(min_row=2):
        r[4].fill = PatternFill("solid", fgColor="FFF4CE")
        for h in r:
            h.alignment = UST

    so = wb.create_sheet("Ortaklık ve Sicil")
    _baslik(so, ["Tür", "Ad / Tarih", "Pay / Görev / Açıklama"], (14, 30, 80))
    for x in s["ortaklar"]:
        so.append(["Ortak" + (" (tüzel)" if x["tuzel"] else ""), x["ad"], f"%{x['pay']:g}"])
    for x in s["yonetim"]:
        so.append(["Yönetim", x["ad"], x["gorev"]])
    for x in sorted(s["degisiklikler"], key=lambda d: d["tarih"] or date.min, reverse=True):
        so.append([f"Değişiklik · {x['tur']}", x["tarih"], x["aciklama"]])
        so.cell(so.max_row, 2).number_format = "DD.MM.YYYY"

    bi = wb.create_sheet("Banka İstihbaratı")
    _baslik(bi, ["Banka", "Nakdi Limit (istihbarat)", "Nakdi Risk (istihbarat)", "Nakdi Risk (KKB)", "Olumsuz İfade", "Metin"], (20, 18, 18, 16, 18, 90))
    kkb = {katla(x["banka"]): x for x in k["kkb"]}
    for x in k["istihbarat"]:
        kx = kkb.get(katla(x["banka"]))
        bi.append([x["banka"], x["nl"] and float(x["nl"]), x["nr"] and float(x["nr"]), kx and float(kx["nr"]), ", ".join(x["olumsuz"]), x["metin"]])
        for c in (2, 3, 4):
            bi.cell(bi.max_row, c).number_format = "#,##0"
        bi.cell(bi.max_row, 6).alignment = UST

    kr = wb.create_sheet("KKB Risk")
    _baslik(kr, ["Banka", "Nakdi Limit", "Nakdi Risk", "Doluluk", "Gayrinakdi Limit", "Gayrinakdi Risk", "İstihbarat Var mı?"], (20, 15, 15, 9, 16, 16, 16))
    ist = {katla(x["banka"]) for x in k["istihbarat"]}
    for x in k["kkb"] + ([{"banka": "TOPLAM", **k["risk"]}] if k["kkb"] else []):
        kr.append([x["banka"], float(x["nl"]), float(x["nr"]), float(x["nr"] / x["nl"]) if x["nl"] else None, float(x["gl"]), float(x["gr"]),
                   "" if x["banka"] == "TOPLAM" else ("Evet" if katla(x["banka"]) in ist else "Hayır")])
        for c in (2, 3, 5, 6):
            kr.cell(kr.max_row, c).number_format = "#,##0"
        kr.cell(kr.max_row, 4).number_format = "0.0%"
        if kr.cell(kr.max_row, 7).value == "Hayır":
            kr.cell(kr.max_row, 7).fill = PatternFill("solid", fgColor=ONEM["orta"])

    ok = wb.create_sheet("Olumsuz Kayıtlar")
    _baslik(ok, ["Tür", "Tarih", "Tutar", "Durum", "Açıklama"], (20, 12, 14, 12, 60))
    for x in k["olumsuz"]:
        ok.append([x["tur"], x["tarih"], x["tutar"] and float(x["tutar"]), x["durum"], x["aciklama"]])
        ok.cell(ok.max_row, 2).number_format = "DD.MM.YYYY"
        ok.cell(ok.max_row, 3).number_format = "#,##0"
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

    md = cikti.with_suffix(".md")
    sat = [f"# Firma İstihbarat Raporu (taslak) — {s['unvan'] or ''}", "", f"Rapor tarihi: {bugun:%d.%m.%Y} · Kaynaklar: " +
           ", ".join(d.ad for d in metinler), "", "## Künye", "", g(y["firma_kunyesi"]), "", "## Ortaklık ve yönetim", "", g(y["ortaklik_ve_yonetim"])]
    if s["ortaklar"]:
        sat += ["", "| Ortak | Pay |", "|---|---|"] + [f"| {x['ad']} | %{x['pay']:g} |" for x in s["ortaklar"]]
    sat += ["", "## Faaliyet ve piyasa", "", g(y["faaliyet_ve_piyasa"]), "", "## Banka ilişkileri", "", g(y["banka_iliskileri"])]
    if k["kkb"]:
        sat += ["", "| Banka | Nakdi limit | Nakdi risk | Gayrinakdi limit | Gayrinakdi risk |", "|---|---:|---:|---:|---:|"]
        sat += [f"| {x['banka']} | {tl(x['nl'])} | {tl(x['nr'])} | {tl(x['gl'])} | {tl(x['gr'])} |" for x in k["kkb"]]
    sat += ["", "## Olumsuz bilgiler", "", g(y["olumsuz_bilgiler"]), "", "## Kod kontrolleri", ""]
    sat += [f"- **{o_}** {a} ({kay})" for o_, a, kay in k["bulgular"]] or ["- Bulgu yok"]
    sat += ["", "## Kaynaklar arası çelişkiler", ""] + ([f"- {g(x['aciklama'])} ({x['kaynak']})" for x in y["celiskiler"]] or ["- Bulunmadı"])
    sat += ["", "## Güçlü yönler", ""] + [f"- {g(x['aciklama'])} ({x['kaynak']})" for x in y["guclu_yonler"]]
    sat += ["", "## Risk işaretleri", ""] + [f"- **{x['onem']}** {g(x['aciklama'])} ({x['kaynak']})" for x in y["risk_isaretleri"]]
    sat += ["", "## Teyit edilecekler", ""] + [f"- {g(t)}" for t in y["teyit_edilecekler"]]
    sat += ["", "## Genel değerlendirme (taslak)", "", g(y["genel_degerlendirme"])]
    if y["dogrulanamayan_sayilar"]:
        sat += ["", f"> Girdilerde bulunmayan sayılar (kontrol edin): {', '.join(y['dogrulanamayan_sayilar'])}"]
    sat += ["", "> Taslaktır. İstihbarat bilgileri bankacılık sırrıdır; yalnız yetkili kişilerle paylaşın."]
    md.write_text("\n".join(sat) + "\n", encoding="utf-8")
    return md


def calistir(klasor: Path, cikti: Path, terimler: list[str] | None = None, bugun: date | None = None, evet: bool = False) -> dict:
    if not klasor.is_dir():
        raise llm.LLMHatasi(f"{klasor} klasörü bulunamadı.")
    bugun = bugun or date.today()
    girdi, metinler, atlanan = dosyalari_oku(klasor)
    if not metinler:
        raise llm.LLMHatasi(f"{klasor}: okunabilir dosya yok.")
    k = kontrol_et(girdi, metinler, bugun)
    harita = takma_adlar(k["sicil"], terimler or [])
    for d in metinler:
        print(f"[OK] {d.ad}: {d.tur}")
    kisi = sum(1 for t in harita.values() if t.startswith("[KİŞİ") and not t.endswith("SOYAD]"))
    llm.onay_al(f"{len(metinler)} dosyanın metni ve kod özeti gönderilecek (firma unvanı, VKN, {kisi} kişi ve "
                f"{sum(1 for t in harita.values() if t.startswith('[ŞİRKET'))} ortak şirket takma adlı; telefon, e-posta, TCKN, IBAN maskeli).", evet)
    y = model_yaz(metinler, k, harita)
    md = rapor_yaz(cikti, metinler, atlanan, k, y, harita, bugun)
    return {"kontrol": k, "yanit": y, "harita": harita, "md": md, "metinler": metinler}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Ticaret sicili, banka istihbaratı, piyasa görüşmeleri ve risk bilgilerinden istihbarat raporu derler.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "ornek_metal", help="Firma istihbarat klasörü")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY (varsayılan: bugün; örnek veride 08.10.2026)")
    p.add_argument("--gizle", nargs="*", default=[], metavar="AD", help="Ayrıca maskelenecek kişi/şirket adları")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "istihbarat_raporu.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    bugun = tarih(a.tarih) if a.tarih else (date(2026, 10, 8) if a.girdi == BURASI / "ornek_veri" / "ornek_metal" else None)
    try:
        s = calistir(a.girdi, a.cikti, a.gizle, bugun, a.evet)
    except llm.LLMHatasi as h:
        print(f"[X] {h}")
        return 1
    for o, ac, kay in s["kontrol"]["bulgular"]:
        if o != "bilgi":
            print(f"[{'X' if o == 'yüksek' else '!'}] {ac}")
    if s["yanit"]["dogrulanamayan_sayilar"]:
        print(f"[!] Girdilerde olmayan sayılar: {', '.join(s['yanit']['dogrulanamayan_sayilar'])}")
    print(f"[OK] Rapor: {a.cikti.resolve()} · {s['md'].resolve()}")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
