"""
Konşimento Bilgi Çıkarma — Gerisi Bende AI Agent
Deniz ve Hava Yolu Operasyon › Denizyolu Operasyon Uzmanı

1. Model konşimento (B/L) metninden yükleyici, alıcı, ihbar adresi, gemi / sefer, yükleme ve boşaltma limanı,
   konteynerler (no, mühür, tip, kap, brüt ağırlık, hacim), navlun ödeme şekli, orijinal adedi ve tarihleri çıkarır.
2. Kod çıkarılan bilgiyi denetler:
   - Her metin alanı ve sayı belgede gerçekten geçiyor mu (uydurma yok)
   - Konteyner numarasının ISO 6346 kontrol hanesi doğru mu
   - Konteyner kap / ağırlık / hacim toplamları belgedeki TOPLAM satırıyla tutuyor mu
   - Aynı konteyner iki kez yazılmış mı
   Sorun varsa modele geri gönderilir (en fazla --tur kez). Kalan sorunlar raporda işaretlenir.
3. İsteğe bağlı çeki listesi (Konteyner No, Kap Adedi, Brüt Kg) verilirse konşimento onunla karşılaştırılır.
   Emre yazılı (TO ORDER) konşimento, "said to contain / shipper's load" kaydı gibi durumlar uyarı olarak eklenir.

Kullanım:
    python agent.py                                   # örnek konşimentoyla dener
    python agent.py --girdi bl.pdf --ceki ceki_listesi.xlsx
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
N = {"type": ["number", "null"]}
KONTEYNER = {"type": "object", "properties": {"no": S, "muhur": S, "tip": S, "kap_adedi": N, "kap_cinsi": S, "tanim": S, "brut_kg": N, "hacim_cbm": N},
             "required": ["no", "muhur", "tip", "kap_adedi", "kap_cinsi", "tanim", "brut_kg", "hacim_cbm"], "additionalProperties": False}
METIN_ALANLARI = ["bl_no", "booking_no", "yukleyici", "alici", "ihbar_adresi", "gemi", "sefer_no", "teslim_alma_yeri", "yukleme_limani",
                  "bosaltma_limani", "teslim_yeri", "duzenleme_yeri", "duzenleme_tarihi", "yukleme_tarihi"]
SEMA = {
    "type": "object",
    "properties": {
        **{a: S for a in METIN_ALANLARI},
        "konteynerler": {"type": "array", "items": KONTEYNER},
        "toplam": {"type": "object", "properties": {"konteyner_adedi": N, "kap_adedi": N, "brut_kg": N, "hacim_cbm": N},
                   "required": ["konteyner_adedi", "kap_adedi", "brut_kg", "hacim_cbm"], "additionalProperties": False},
        "navlun": {"type": "string", "enum": ["Prepaid (peşin)", "Collect (varışta)", "Belirtilmemiş"]},
        "orijinal_adedi": N,
        "gtip": {"type": "array", "items": S},
        "kayitlar": {"type": "array", "items": S},
    },
    "required": [*METIN_ALANLARI, "konteynerler", "toplam", "navlun", "orijinal_adedi", "gtip", "kayitlar"],
    "additionalProperties": False,
}
ETIKET = {"bl_no": "B/L No", "booking_no": "Booking No", "yukleyici": "Yükleyici (Shipper)", "alici": "Alıcı (Consignee)",
          "ihbar_adresi": "İhbar Adresi (Notify)", "gemi": "Gemi", "sefer_no": "Sefer No", "teslim_alma_yeri": "Teslim Alma Yeri",
          "yukleme_limani": "Yükleme Limanı (POL)", "bosaltma_limani": "Boşaltma Limanı (POD)", "teslim_yeri": "Teslim Yeri",
          "duzenleme_yeri": "Düzenleme Yeri", "duzenleme_tarihi": "Düzenleme Tarihi", "yukleme_tarihi": "Gemiye Yükleme Tarihi"}

_HARF = {}
_v = 10
for _c in "ABCDEFGHIJKLMNOPQRSTUVWXYZ":
    if _v % 11 == 0:
        _v += 1
    _HARF[_c] = _v
    _v += 1


def iso6346(no: str) -> str | None:
    """Geçerliyse None; değilse sorun açıklaması."""
    n = re.sub(r"[\s-]", "", str(no or "")).upper()
    if not re.fullmatch(r"[A-Z]{3}[UJZ]\d{7}", n):
        return "biçim 4 harf (son harf U/J/Z) + 7 rakam olmalı"
    hane = sum((_HARF[c] if c.isalpha() else int(c)) * 2 ** i for i, c in enumerate(n[:10])) % 11 % 10
    return None if hane == int(n[10]) else f"kontrol hanesi {n[10]}, olması gereken {hane}"


def sayi_belgede(v, metin: str) -> bool:
    return v is None or sayi_yaz(v) in sayilar(metin)


def denetle(t: dict, metin: str) -> list[str]:
    sorun = []
    for a in METIN_ALANLARI:
        v = t[a].strip()
        if v and not metinde_var(v, metin) and not all(metinde_var(p, metin) for p in re.split(r"[,\n]", v) if p.strip()):
            sorun.append(f"{ETIKET[a]}: '{v[:60]}' belgede geçmiyor (belgedeki yazımı aynen kullanın)")
    gorulen = set()
    for k in t["konteynerler"]:
        no = re.sub(r"[\s-]", "", k["no"]).upper()
        if no in gorulen:
            sorun.append(f"Konteyner {no} iki kez yazılmış")
        gorulen.add(no)
        if not metinde_var(no, metin):
            sorun.append(f"Konteyner {no} belgede geçmiyor")
        elif (h := iso6346(no)):
            sorun.append(f"Konteyner {no}: ISO 6346 {h} — belgeden yeniden okuyun; belgede de böyleyse not edin")
        for alan, ad in (("kap_adedi", "kap"), ("brut_kg", "brüt kg"), ("hacim_cbm", "hacim")):
            if not sayi_belgede(k[alan], metin):
                sorun.append(f"Konteyner {no}: {ad} {k[alan]} belgede geçmiyor")
        if k["muhur"] and not metinde_var(k["muhur"], metin):
            sorun.append(f"Konteyner {no}: mühür {k['muhur']} belgede geçmiyor")
    top = t["toplam"]
    for alan, ad in (("kap_adedi", "kap"), ("brut_kg", "brüt kg"), ("hacim_cbm", "hacim (CBM)")):
        if not sayi_belgede(top[alan], metin):
            sorun.append(f"Toplam {ad} {top[alan]} belgede geçmiyor")
        degerler = [k[alan] for k in t["konteynerler"]]
        if top[alan] is not None and degerler and None not in degerler and abs(sum(degerler) - top[alan]) > 0.01:
            sorun.append(f"Konteynerlerin {ad} toplamı {sum(degerler):g}, TOPLAM satırı {top[alan]:g} — belgeyi yeniden okuyun")
    if top["konteyner_adedi"] is not None and top["konteyner_adedi"] != len(t["konteynerler"]):
        sorun.append(f"Toplam konteyner adedi {top['konteyner_adedi']:g}, listelenen {len(t['konteynerler'])}")
    for g in t["gtip"]:
        if not metinde_var(g, metin):
            sorun.append(f"GTİP / HS {g} belgede geçmiyor")
    return sorun


def uyarilar(t: dict, metin: str) -> list[str]:
    """Kodun belgeden ve çıkarılan bilgiden ürettiği uyarılar (model sonucu değiştirmez)."""
    u, k = [], katla(metin)
    if "to order" in katla(t["alici"]) or "emrine" in katla(t["alici"]):
        u.append("Emre yazılı konşimento: malın teslimi için orijinal konşimentonun ciro zinciri ile ibrazı gerekir")
    if "said to contain" in k or "shipper s load" in k or "stc" in k.split():
        u.append("'Shipper's load, stow and count / said to contain' kaydı: taşıyan, konteyner içeriğini saymadığını belirtiyor")
    if t["navlun"] == "Belirtilmemiş":
        u.append("Navlun ödeme şekli (prepaid / collect) belgede bulunamadı")
    if t["orijinal_adedi"] is None:
        u.append("Orijinal konşimento adedi bulunamadı")
    if "telex release" in k or "surrender" in k:
        u.append("Telex release / surrender kaydı var: teslim orijinal ibrazı olmadan yapılabilir, acenteyle teyit edin")
    if "seaway bill" in k or "sea waybill" in k:
        u.append("Belge sea waybill olabilir: kıymetli evrak değildir, ciro edilemez")
    yt, dt = tarih_coz(t["yukleme_tarihi"]), tarih_coz(t["duzenleme_tarihi"])
    if yt and dt and dt < yt:
        u.append(f"Düzenleme tarihi ({dt:%d.%m.%Y}) gemiye yükleme tarihinden ({yt:%d.%m.%Y}) önce")
    return u


def ceki_karsilastir(t: dict, ceki: list[dict]) -> list[list]:
    bl = {re.sub(r"[\s-]", "", k["no"]).upper(): k for k in t["konteynerler"]}
    out = []
    for c in ceki:
        no = re.sub(r"[\s-]", "", str(c["no"])).upper()
        k = bl.pop(no, None)
        if not k:
            out.append([no, "Çeki listesinde var, konşimentoda yok", c["kap"], "", c["kg"], ""])
            continue
        for alan, ad, cv in (("kap_adedi", "kap", c["kap"]), ("brut_kg", "brüt kg", c["kg"])):
            try:
                cv = float(sayi_norm(cv)) if cv not in ("", None) else None
            except ValueError:
                cv = None
            if cv is not None and k[alan] is not None and abs(cv - k[alan]) > 0.01:
                out.append([no, f"{ad} farklı", cv if alan == "kap_adedi" else "", k[alan] if alan == "kap_adedi" else "",
                            cv if alan == "brut_kg" else "", k[alan] if alan == "brut_kg" else ""])
    for no in bl:
        out.append([no, "Konşimentoda var, çeki listesinde yok", "", bl[no]["kap_adedi"], "", bl[no]["brut_kg"]])
    return out


def uret(metin: str, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    return llm.json_iste(sistem, f"<konsimento>\n{llm.maskele(metin)}\n</konsimento>" + duzeltme_blogu(duzeltme, onceki), SEMA)


def rapor_yaz(cikti: Path, t: dict, kalan: list[str], tur: int, uyari: list[str], fark: list[list] | None) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Konşimento"
    baslik(ws, ["Alan", "Değer", "Kontrol"], (26, 70, 14))
    for a in METIN_ALANLARI:
        ws.append([ETIKET[a], t[a], "Belgede var" if not any(s.startswith(ETIKET[a]) for s in kalan) else "KONTROL"])
        boya(ws, [2], AI_DOLGU)
    for ad, v in (("Navlun", t["navlun"]), ("Orijinal adedi", t["orijinal_adedi"]), ("GTİP / HS", ", ".join(t["gtip"])),
                  ("Kayıtlar (clause)", "; ".join(t["kayitlar"]))):
        ws.append([ad, v, ""])
        boya(ws, [2], AI_DOLGU)
    if uyari:
        ws.append([])
        for x in uyari:
            ws.append(["Uyarı", x])
            boya(ws, [1, 2], SARI)
    son_not(ws, "Önemli: Yapay zekâ ile çıkarılmıştır. Gümrük beyanı ve akreditif ibrazından önce orijinal konşimentoyla karşılaştırın.", kalan, tur)
    hizala(ws)

    k = wb.create_sheet("Konteynerler")
    baslik(k, ["Konteyner No", "ISO 6346", "Mühür", "Tip", "Kap Adedi", "Kap Cinsi", "Tanım", "Brüt Kg", "Hacim (CBM)"], (16, 30, 14, 8, 10, 12, 36, 12, 12))
    for x in t["konteynerler"]:
        h = iso6346(x["no"])
        k.append([x["no"], h or "Geçerli", x["muhur"], x["tip"], x["kap_adedi"], x["kap_cinsi"], x["tanim"], x["brut_kg"], x["hacim_cbm"]])
        boya(k, [2], KIRMIZI if h else YESIL)
    top = t["toplam"]
    k.append(["TOPLAM (belge)", "", "", "", top["kap_adedi"], "", f"{top['konteyner_adedi'] or ''} konteyner", top["brut_kg"], top["hacim_cbm"]])
    for h in k[k.max_row]:
        h.font = Font(bold=True)
    for r in range(2, k.max_row + 1):
        k.cell(r, 8).number_format = "#,##0.00"
        k.cell(r, 9).number_format = "#,##0.00"
    k.freeze_panes = "A2"

    if fark is not None:
        f = wb.create_sheet("Çeki Karşılaştırma")
        baslik(f, ["Konteyner No", "Durum", "Çeki Kap", "B/L Kap", "Çeki Kg", "B/L Kg"], (16, 40, 10, 10, 12, 12))
        for r in fark:
            f.append(r)
            boya(f, [2], KIRMIZI)
        if not fark:
            f.append(["—", "Çeki listesi ile konşimento uyumlu"])
            boya(f, [2], YESIL)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(girdi: Path, cikti: Path, ceki_yolu: Path | None = None, tur: int = 2, evet: bool = False) -> dict:
    if not girdi.exists():
        raise belge.BelgeHatasi(f"{girdi}: dosya bulunamadı.")
    metin = belge.metin_oku(girdi)
    if len(metin.strip()) < 50:
        raise belge.BelgeHatasi(f"{girdi.name}: metin bulunamadı.")
    ceki = None
    if ceki_yolu:
        ceki = [c for c in kayitlar(ceki_yolu, {"no": ("Konteyner No", "Konteyner", "Container No", "Container"), "kap": ("Kap Adedi", "Kap", "Packages", "Koli"),
                                                 "kg": ("Brüt Kg", "Brüt", "Gross Weight", "Brut Kg")}, ("no",)) if c["no"]]
    llm.onay_al(f"Konşimento metni ({len(metin)} karakter; telefon / e-posta maskeli) bilgi çıkarma için gönderilecek.", evet)
    t, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(metin, d, o), lambda x: denetle(x, metin), tur)
    uyari = uyarilar(t, metin)
    fark = ceki_karsilastir(t, ceki) if ceki is not None else None
    rapor_yaz(cikti, t, kalan, yapilan, uyari, fark)
    js = cikti.with_suffix(".json")
    js.write_text(json.dumps(t, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"taslak": t, "kalan": kalan, "tur": yapilan, "uyari": uyari, "fark": fark, "json": js}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    p = argparse.ArgumentParser(description="Konşimentodan (B/L) yükleyici, alıcı, liman, konteyner ve yük bilgilerini çıkarır ve denetler.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "konsimento.txt", help="Konşimento (.pdf, .docx, .txt)")
    p.add_argument("--ceki", type=Path, help="İsteğe bağlı çeki listesi (.xlsx/.csv): Konteyner No, Kap Adedi, Brüt Kg")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "konsimento.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    if a.girdi == BURASI / "ornek_veri" / "konsimento.txt" and not a.ceki:
        a.ceki = BURASI / "ornek_veri" / "ceki_listesi.csv"
    try:
        s = calistir(a.girdi, a.cikti, a.ceki, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    t = s["taslak"]
    print(f"[OK] B/L {t['bl_no']} · {len(t['konteynerler'])} konteyner · {t['yukleme_limani']} → {t['bosaltma_limani']}")
    for x in s["uyari"] + [f"Kalan: {k}" for k in s["kalan"]] + [f"Çeki: {r[0]} {r[1]}" for r in s["fark"] or []]:
        print(f"[!] {x}")
    print(f"[OK] Rapor: {a.cikti.resolve()} (+ {s['json'].name})")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
