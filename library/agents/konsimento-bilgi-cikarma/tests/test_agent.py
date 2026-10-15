"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri" / "konsimento.txt"
CEKI = KLASOR / "ornek_veri" / "ceki_listesi.csv"
METIN = ORNEK.read_text(encoding="utf-8")
MESAJLAR = []
DOGRU = {
    "bl_no": "WWLU-IST-2610-0457", "booking_no": "88124570",
    "yukleyici": "Örnek Tekstil Sanayi ve Ticaret A.Ş., Organize Sanayi Bölgesi 5. Cadde No: 12, Bursa, TÜRKİYE",
    "alici": "TO ORDER OF EXAMPLE BANK GMBH", "ihbar_adresi": "Muster Handels GmbH, Hafenstrasse 20, 20457 Hamburg, GERMANY",
    "gemi": "MV EXAMPLE STAR", "sefer_no": "642W", "teslim_alma_yeri": "BURSA, TURKEY", "yukleme_limani": "GEMLIK, TURKEY (TRGEM)",
    "bosaltma_limani": "HAMBURG, GERMANY (DEHAM)", "teslim_yeri": "HAMBURG CY", "duzenleme_yeri": "ISTANBUL", "duzenleme_tarihi": "14.10.2026",
    "yukleme_tarihi": "14.10.2026",
    "konteynerler": [
        {"no": "WWKU1234564", "muhur": "TR0457881", "tip": "40HC", "kap_adedi": 620, "kap_cinsi": "CARTONS", "tanim": "100% COTTON KNITTED T-SHIRTS", "brut_kg": 8450, "hacim_cbm": 67.5},
        {"no": "WWKU7654327", "muhur": "TR0457882", "tip": "40HC", "kap_adedi": 580, "kap_cinsi": "CARTONS", "tanim": "100% COTTON KNITTED T-SHIRTS", "brut_kg": 7920.5, "hacim_cbm": 66},
        {"no": "WWKU2401178", "muhur": "TR0457883", "tip": "20GP", "kap_adedi": 210, "kap_cinsi": "CARTONS", "tanim": "COTTON SWEATSHIRTS", "brut_kg": 3105, "hacim_cbm": 28.2}],
    "toplam": {"konteyner_adedi": 3, "kap_adedi": 1410, "brut_kg": 19475.5, "hacim_cbm": 161.7},
    "navlun": "Prepaid (peşin)", "orijinal_adedi": 3, "gtip": ["6109.10", "6110.20"], "kayitlar": ["SHIPPER'S LOAD, STOW AND COUNT", "SAID TO CONTAIN"]}


def hatali() -> dict:
    t = copy.deepcopy(DOGRU)
    t["gemi"] = "MV EXAMPLE STAR II"
    t["konteynerler"][1]["kap_adedi"] = 850
    return t


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return copy.deepcopy(DOGRU) if "<duzeltme>" in kullanici else hatali()


class KodTesti(unittest.TestCase):
    def test_iso6346(self):
        self.assertIsNone(agent.iso6346("CSQU3054383"))
        self.assertIsNone(agent.iso6346("WWKU 123456-4"))
        self.assertIn("olması gereken 3", agent.iso6346("CSQU3054384"))
        self.assertIn("biçim", agent.iso6346("ABC1234567"))

    def test_denetim(self):
        self.assertEqual(agent.denetle(DOGRU, METIN), [])
        s = agent.denetle(hatali(), METIN)
        self.assertEqual(len(s), 3)
        self.assertTrue(any("Gemi" in x for x in s) and any("kap 850" in x for x in s) and any("toplamı 1680" in x for x in s))
        t = copy.deepcopy(DOGRU)
        t["konteynerler"][0]["no"] = "WWKU1234565"
        t["konteynerler"].append(copy.deepcopy(t["konteynerler"][2]))
        s = agent.denetle(t, METIN)
        self.assertTrue(any("WWKU1234565 belgede geçmiyor" in x for x in s) and any("iki kez" in x for x in s) and any("adedi 3" in x for x in s))

    def test_uyarilar_ve_ceki(self):
        u = agent.uyarilar(DOGRU, METIN)
        self.assertEqual(len(u), 2)
        self.assertTrue(u[0].startswith("Emre yazılı"))
        ceki = [{"no": "WWKU1234564", "kap": "620", "kg": "8450"}, {"no": "WWKU2401178", "kap": "210", "kg": "3150"}, {"no": "WWKU9999990", "kap": "1", "kg": ""}]
        f = agent.ceki_karsilastir(DOGRU, ceki)
        self.assertEqual([(r[0], r[1]) for r in f], [("WWKU2401178", "brüt kg farklı"), ("WWKU9999990", "Çeki listesinde var, konşimentoda yok"),
                                                     ("WWKU7654327", "Konşimentoda var, çeki listesinde yok")])


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "k.xlsx"
            s = agent.calistir(ORNEK, cikti, CEKI, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            self.assertIn("toplamı 1680", MESAJLAR[1].split("<duzeltme>")[1])
            self.assertEqual([r[0] for r in s["fark"]], ["WWKU2401178"])
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Konşimento", "Konteynerler", "Çeki Karşılaştırma"])
            self.assertEqual(wb["Konteynerler"].cell(2, 2).value, "Geçerli")
            self.assertIn('"bl_no": "WWLU-IST-2610-0457"', s["json"].read_text(encoding="utf-8"))

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(ORNEK, Path(tmp) / "k.xlsx", tur=0, evet=True)
            self.assertEqual((len(s["kalan"]), s["fark"]), (3, None))
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--girdi", str(ORNEK.with_name("yok.txt")), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
