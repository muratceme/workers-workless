"""
Karar ve İçtihat Özeti — Workers / Workless AI Agent
Dava Takip › Avukat

1. Kod karar metnini satır (paragraf) satır numaralar.
2. Model kararı başlıklarla özetler: mahkeme / daire, esas ve karar no, karar tarihi, taraflar, olay, tarafların iddia ve
   savunmaları, yargılama süreci, hukuki mesele, gerekçe, hüküm, oy durumu, atıf yapılan mevzuat ve kararlar,
   anahtar kelimeler. Her bölüm dayandığı satırları gösterir; hüküm için metinden aynen alıntı verir.
3. Kod denetler:
   - Esas / karar no ve karar tarihi metinde geçiyor mu
   - Mevzuat maddesinin numarası dayandığı satırda yazıyor mu
   - Atıf yapılan kararların esas / karar numaraları metinde var mı
   - Hüküm alıntısı metinde var mı
   - Oy durumu metindeki "oybirliği / oyçokluğu" ifadesiyle tutarlı mı
   - Dayanak satırları geçerli mi
   Sorun varsa modele geri gönderilir (en fazla --tur kez).

Kullanım:
    python agent.py                                      # örnek (kurgusal) karar metniyle dener
    python agent.py --girdi karar.pdf --tur 3
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
K = {"type": "array", "items": {"type": "integer"}}


def _b(**alan) -> dict:
    alan["kaynak_satirlar"] = K
    return {"type": "object", "properties": alan, "required": list(alan), "additionalProperties": False}


def _l(**alan) -> dict:
    return {"type": "array", "items": _b(**alan)}


SEMA = {
    "type": "object",
    "properties": {
        "kunye": _b(mahkeme=S, esas_no=S, karar_no=S, karar_tarihi=S),
        "taraflar": _l(rol=S, ad=S),
        "olay": _b(metin=S),
        "iddia_savunma": _b(metin=S),
        "yargilama_sureci": _b(metin=S),
        "hukuki_mesele": _b(metin=S),
        "gerekce": _b(metin=S),
        "hukum": _b(metin=S, alinti=S, sonuc={"type": "string", "enum": ["Onama", "Bozma", "Kısmen bozma", "Kabul", "Kısmen kabul", "Ret", "Düzelterek onama", "Diğer"]}),
        "oy_durumu": {"type": "string", "enum": ["Oybirliği", "Oyçokluğu", "Belirtilmemiş"]},
        "mevzuat": _l(kanun=S, madde=S, konu=S),
        "atif_kararlar": _l(mahkeme=S, esas_no=S, karar_no=S, ozet=S),
        "anahtar_kelimeler": {"type": "array", "items": S},
    },
    "required": ["kunye", "taraflar", "olay", "iddia_savunma", "yargilama_sureci", "hukuki_mesele", "gerekce", "hukum", "oy_durumu", "mevzuat", "atif_kararlar",
                 "anahtar_kelimeler"],
    "additionalProperties": False,
}
BOLUMLER = (("olay", "Olay"), ("iddia_savunma", "İddia ve savunma"), ("yargilama_sureci", "Yargılama süreci"), ("hukuki_mesele", "Hukuki mesele"),
            ("gerekce", "Gerekçe"), ("hukum", "Hüküm"))


def satirlar(metin: str) -> list[str]:
    return [s.strip() for s in metin.splitlines() if s.strip()]


def denetle(t: dict, l: list[str]) -> list[str]:
    n, metin, sorun = len(l), "\n".join(l), []

    def kaynak(ad: str, ks: list[int]) -> list[str]:
        if not ks:
            sorun.append(f"{ad}: kaynak satır numarası yok")
        gec = [k for k in ks if not 1 <= k <= n]
        if gec:
            sorun.append(f"{ad}: metinde olmayan satır {gec}")
        return [l[k - 1] for k in ks if 1 <= k <= n]

    k = t["kunye"]
    kaynak("Künye", k["kaynak_satirlar"])
    for alan, ad in (("esas_no", "Esas no"), ("karar_no", "Karar no")):
        if k[alan] and not metinde_var(k[alan], metin):
            sorun.append(f"Künye: {ad} '{k[alan]}' metinde geçmiyor")
    if k["karar_tarihi"] and not metinde_var(k["karar_tarihi"], metin):
        sorun.append(f"Künye: karar tarihi '{k['karar_tarihi']}' metinde geçmiyor")
    for x in t["taraflar"]:
        kaynak(f"Taraf {x['rol']}", x["kaynak_satirlar"])
    for anahtar, ad in BOLUMLER:
        kaynak(ad, t[anahtar]["kaynak_satirlar"])
    if not metinde_var(t["hukum"]["alinti"], metin):
        sorun.append(f"Hüküm: alıntı metinde yok: '{t['hukum']['alinti'][:60]}'")
    for m in t["mevzuat"]:
        sat = kaynak(f"Mevzuat {m['kanun']} m.{m['madde']}", m["kaynak_satirlar"])
        no = re.findall(r"\d+", m["madde"])
        if not no or no[0] not in {x for s in sat for x in re.findall(r"\d+", s)}:
            sorun.append(f"Mevzuat {m['kanun']} m.{m['madde']}: madde numarası dayandığı satırda yazmıyor")
    for a in t["atif_kararlar"]:
        kaynak(f"Atıf {a['esas_no']}", a["kaynak_satirlar"])
        for alan in ("esas_no", "karar_no"):
            if a[alan] and not metinde_var(a[alan], metin):
                sorun.append(f"Atıf yapılan karar: {alan} '{a[alan]}' metinde geçmiyor")
    km = katla(metin)
    beklenen = "Oybirliği" if "oybirlig" in km else ("Oyçokluğu" if "oycoklug" in km else None)
    if beklenen and t["oy_durumu"] != beklenen:
        sorun.append(f"Oy durumu: metinde '{beklenen.lower()}' yazıyor, '{t['oy_durumu']}' verilmiş")
    if not beklenen and t["oy_durumu"] != "Belirtilmemiş":
        sorun.append(f"Oy durumu: metinde oybirliği / oyçokluğu ifadesi yok, '{t['oy_durumu']}' verilmiş")
    return sorun


def uret(l: list[str], duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    numarali = "\n".join(f"{i}: {llm.maskele(s)}" for i, s in enumerate(l, 1))
    return llm.json_iste(sistem, f"<karar>\n{numarali}\n</karar>" + duzeltme_blogu(duzeltme, onceki), SEMA)


def ozet_metni(t: dict) -> str:
    k = t["kunye"]
    out = [f"# {k['mahkeme']} · E. {k['esas_no']} · K. {k['karar_no']} · {k['karar_tarihi']}", "",
           "**Taraflar:** " + "; ".join(f"{x['rol']}: {x['ad']}" for x in t["taraflar"]), ""]
    for anahtar, ad in BOLUMLER:
        out += [f"## {ad}", t[anahtar]["metin"] + f" _(satır {', '.join(map(str, t[anahtar]['kaynak_satirlar']))})_", ""]
    out += [f"> {t['hukum']['alinti']}", "", f"**Sonuç:** {t['hukum']['sonuc']} · **Oy durumu:** {t['oy_durumu']}", ""]
    if t["mevzuat"]:
        out += ["## Mevzuat"] + [f"- {m['kanun']} m. {m['madde']}: {m['konu']}" for m in t["mevzuat"]] + [""]
    if t["atif_kararlar"]:
        out += ["## Atıf yapılan kararlar"] + [f"- {a['mahkeme']} E. {a['esas_no']} K. {a['karar_no']}: {a['ozet']}" for a in t["atif_kararlar"]] + [""]
    out.append("**Anahtar kelimeler:** " + ", ".join(t["anahtar_kelimeler"]))
    return "\n".join(out)


def rapor_yaz(cikti: Path, l, t, kalan, tur) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Karar Özeti"
    k = t["kunye"]
    baslik(ws, ["Başlık", "İçerik", "Dayanak Satırlar"], (22, 100, 14))
    for ad, v, ks in (("Mahkeme / Daire", k["mahkeme"], k["kaynak_satirlar"]), ("Esas No", k["esas_no"], []), ("Karar No", k["karar_no"], []),
                      ("Karar Tarihi", k["karar_tarihi"], []), ("Taraflar", "; ".join(f"{x['rol']}: {x['ad']}" for x in t["taraflar"]), [])):
        ws.append([ad, v, ", ".join(map(str, ks))])
    for anahtar, ad in BOLUMLER:
        ws.append([ad, t[anahtar]["metin"], ", ".join(map(str, t[anahtar]["kaynak_satirlar"]))])
        boya(ws, [2], AI_DOLGU)
    ws.append(["Hüküm (alıntı)", t["hukum"]["alinti"], ""])
    ws.append(["Sonuç / Oy", f"{t['hukum']['sonuc']} · {t['oy_durumu']}", ""])
    ws.append(["Anahtar Kelimeler", ", ".join(t["anahtar_kelimeler"]), ""])
    son_not(ws, "Önemli: Yapay zekâ özetidir. Kullanmadan önce kararın tam metniyle ve resmî kaynakla karşılaştırın; kararın kesinleşme durumunu ayrıca kontrol edin.",
            kalan, tur)
    hizala(ws)
    m = wb.create_sheet("Mevzuat ve Atıflar")
    baslik(m, ["Tür", "Kaynak", "Madde / Esas", "Karar No", "Konu / Özet", "Dayanak Satırlar"], (10, 30, 14, 14, 60, 14))
    for x in t["mevzuat"]:
        m.append(["Mevzuat", x["kanun"], x["madde"], "", x["konu"], ", ".join(map(str, x["kaynak_satirlar"]))])
    for x in t["atif_kararlar"]:
        m.append(["Karar", x["mahkeme"], x["esas_no"], x["karar_no"], x["ozet"], ", ".join(map(str, x["kaynak_satirlar"]))])
    hizala(m)
    d = wb.create_sheet("Metin")
    baslik(d, ["Satır", "Metin"], (6, 120))
    for i, s in enumerate(l, 1):
        d.append([i, s])
    hizala(d)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    md = cikti.with_suffix(".md")
    md.write_text(ozet_metni(t) + "\n", encoding="utf-8")
    return md


def calistir(girdi: Path, cikti: Path, tur: int = 2, evet: bool = False) -> dict:
    if not girdi.exists():
        raise belge.BelgeHatasi(f"{girdi}: dosya bulunamadı.")
    l = satirlar(belge.metin_oku(girdi))
    if not l:
        raise belge.BelgeHatasi(f"{girdi.name}: metin bulunamadı.")
    print(f"[OK] {len(l)} satır / paragraf")
    llm.onay_al("Karar metni (numaralı satırlar; telefon / e-posta / TCKN maskeli) gönderilecek. Taraf adları maskelenmez.", evet)
    t, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(l, d, o), lambda x: denetle(x, l), tur)
    md = rapor_yaz(cikti, l, t, kalan, yapilan)
    return {"satirlar": l, "taslak": t, "kalan": kalan, "tur": yapilan, "md": md}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    p = argparse.ArgumentParser(description="Mahkeme kararını taraflar, olay, hukuki mesele, gerekçe ve hüküm başlıklarıyla özetler.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "karar.txt", help="Karar metni (.txt, .md, .docx, .pdf)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "karar_ozeti.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.cikti, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    k = s["taslak"]["kunye"]
    print(f"[OK] {k['mahkeme']} E. {k['esas_no']} K. {k['karar_no']} · {s['taslak']['hukum']['sonuc']} · {s['taslak']['oy_durumu']}")
    for x in s["kalan"]:
        print(f"[!] Kalan: {x}")
    print(f"[OK] Rapor: {a.cikti.resolve()} (+ {s['md'].name})")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
