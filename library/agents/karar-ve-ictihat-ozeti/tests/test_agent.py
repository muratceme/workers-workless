"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri" / "karar.txt"
MESAJLAR = []
DOGRU = {
    "kunye": {"mahkeme": "Yargıtay 4. Hukuk Dairesi", "esas_no": "2025/1234", "karar_no": "2026/5678", "karar_tarihi": "12.03.2026", "kaynak_satirlar": [2, 3]},
    "taraflar": [{"rol": "Davacı", "ad": "A.B.", "kaynak_satirlar": [4]}, {"rol": "Davalı", "ad": "Örnek Market Ticaret A.Ş.", "kaynak_satirlar": [4]}],
    "olay": {"metin": "Davacı, davalının marketinde ıslak zeminde düşerek bileğini kırmıştır.", "kaynak_satirlar": [5, 6]},
    "iddia_savunma": {"metin": "Davacı uyarı levhası konulmadığını; davalı zeminin kuru olduğunu ve davacının dikkatsiz olduğunu ileri sürmüştür.", "kaynak_satirlar": [6, 7]},
    "yargilama_sureci": {"metin": "İlk derece kısmen kabul etmiş, istinaf başvurusu esastan reddedilmiş, davalı temyiz etmiştir.", "kaynak_satirlar": [8, 9]},
    "hukuki_mesele": {"metin": "Uyarı levhası koymayan işletmenin sorumluluğu ve zarar görenin müterafik kusuru.", "kaynak_satirlar": [10, 11, 12]},
    "gerekce": {"metin": "Uyarı levhası konulmaması özen yükümlülüğüne aykırıdır; %30 müterafik kusur indirimi dosyaya uygundur.", "kaynak_satirlar": [12, 13, 14]},
    "hukum": {"metin": "Temyiz itirazları reddedilerek karar onanmıştır.", "alinti": "kararın ONANMASINA", "sonuc": "Onama", "kaynak_satirlar": [15]},
    "oy_durumu": "Oybirliği",
    "mevzuat": [{"kanun": "6098 sayılı TBK", "madde": "49", "konu": "Haksız fiil sorumluluğu", "kaynak_satirlar": [10]},
                {"kanun": "6098 sayılı TBK", "madde": "52", "konu": "İndirim sebepleri", "kaynak_satirlar": [11]},
                {"kanun": "6100 sayılı HMK", "madde": "266", "konu": "Bilirkişiye başvurma", "kaynak_satirlar": [13]}],
    "atif_kararlar": [{"mahkeme": "Yargıtay 4. HD", "esas_no": "2024/987", "karar_no": "2024/3456", "ozet": "İşletme sahibinin özen yükümlülüğü", "kaynak_satirlar": [14]}],
    "anahtar_kelimeler": ["haksız fiil", "müterafik kusur", "özen yükümlülüğü", "ıslak zemin"],
}


def hatali() -> dict:
    t = copy.deepcopy(DOGRU)
    t["kunye"]["karar_no"] = "2026/5687"
    t["mevzuat"][1]["kaynak_satirlar"] = [10]
    t["oy_durumu"] = "Oyçokluğu"
    return t


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    return copy.deepcopy(DOGRU) if "<duzeltme>" in kullanici else hatali()


class KodTesti(unittest.TestCase):
    def test_denetim(self):
        l = agent.satirlar(ORNEK.read_text(encoding="utf-8"))
        self.assertEqual(len(l), 15)
        self.assertEqual(agent.denetle(DOGRU, l), [])
        s = agent.denetle(hatali(), l)
        self.assertEqual(len(s), 3)
        self.assertTrue(any("2026/5687" in x for x in s) and any("m.52: madde" in x for x in s) and any("Oy durumu" in x for x in s))
        t = copy.deepcopy(DOGRU)
        t["hukum"]["alinti"] = "kararın BOZULMASINA"
        t["atif_kararlar"][0]["esas_no"] = "2023/111"
        t["olay"]["kaynak_satirlar"] = []
        self.assertEqual(len(agent.denetle(t, l)), 3)
        t = copy.deepcopy(DOGRU)
        t["oy_durumu"] = "Belirtilmemiş"
        self.assertEqual(len(agent.denetle(t, l[:-1])), 2)  # son satır yoksa hüküm alıntısı da yok; oy durumu belirtilmemiş doğru


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "k.xlsx"
            s = agent.calistir(ORNEK, cikti, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (2, [], 1))
            self.assertIn("3: ESAS NO: 2025/1234", MESAJLAR[0])
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Karar Özeti", "Mevzuat ve Atıflar", "Metin"])
            self.assertEqual(wb["Mevzuat ve Atıflar"].max_row, 5)
            md = s["md"].read_text(encoding="utf-8")
            self.assertIn("E. 2025/1234 · K. 2026/5678", md)
            self.assertIn("> kararın ONANMASINA", md)

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(ORNEK, Path(tmp) / "k.xlsx", tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 3)
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--girdi", str(ORNEK.with_name("yok.txt")), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
