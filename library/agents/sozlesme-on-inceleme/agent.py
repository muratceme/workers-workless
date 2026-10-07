"""
Sözleşme Ön İnceleme — Workers / Workless AI Agent
Hukuk › Avukat

1. Sözleşme (.docx, .pdf, .txt) okunur ve kodla maddelere bölünür; her madde konu anahtar kelimeleriyle
   etiketlenir (yapay zekâ yok). IBAN, TCKN, telefon, e-posta ve --gizle ile verilen adlar maskelenir.
2. Maddeler, sizin tarafınızın bakış açısıyla modele gönderilir. Model her madde için risk düzeyi, kimin
   lehine olduğu, gerekçe, müzakere önerisi ve alternatif madde metni yazar; 20 konuluk standart kontrol
   listesinde eksik ve belirsiz düzenlemeleri işaretler.
3. Sonuç Excel'e (ve kısa bir Markdown özetine) yazılır; her öneri için 'Avukat Onayı' sütunu vardır.

Bu bir ön incelemedir, hukuki görüş değildir. Sonuçlar avukat tarafından değerlendirilmelidir.

Kullanım:
    python agent.py                                               # örnek sözleşmeyle dener
    python agent.py --girdi sozlesme.docx --taraf "Müşteri"
    python agent.py --girdi sozlesme.pdf --taraf "Alıcı" --gizle "Ahmet Yılmaz" "Örnek A.Ş."
"""
from __future__ import annotations

import argparse
import re
import sys
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
PAKET_KARAKTER = 60000          # bir istekte gönderilen en fazla sözleşme metni

# Kontrol listesi: (konu, anahtar kelime kökleri — kelime başında aranır)
KONULAR = [
    ("Taraflar ve tanımlar", ("taraflar", "tanımlar", "bundan sonra")),
    ("Konu ve kapsam", ("konusu", "kapsam")),
    ("Süre ve yenileme", ("süre", "yenilen", "kendiliğinden uzar", "uzar")),
    ("Fesih", ("fesih", "feshed", "sona er")),
    ("Ödeme şartları", ("ödeme", "bedel", "fatura", "ücret")),
    ("Gecikme faizi / temerrüt", ("gecikme", "temerrüt", "faiz")),
    ("Ceza şartı", ("cezai şart", "ceza şart", "cezası", "cezai")),
    ("Sorumluluk ve tazminat", ("sorumlu", "tazmin", "zarar")),
    ("Teminat", ("teminat", "kefil", "garanti mektub")),
    ("Gizlilik", ("gizli",)),
    ("Kişisel verilerin korunması (KVKK)", ("kişisel veri", "kvkk", "6698")),
    ("Fikri mülkiyet", ("fikri", "telif", "lisans", "mülkiyet")),
    ("Mücbir sebep", ("mücbir", "force majeure")),
    ("Devir ve temlik", ("devred", "devir", "temlik")),
    ("Alt yüklenici", ("alt yüklenici", "taşeron", "üçüncü kişi")),
    ("Rekabet yasağı / münhasırlık", ("rekabet", "münhasır")),
    ("Uyuşmazlık çözümü (yetki, tahkim, arabuluculuk)", ("mahkeme", "tahkim", "arabulucu", "uyuşmazlık", "icra daire")),
    ("Uygulanacak hukuk", ("uygulanacak hukuk", "türk hukuku", "hukukuna tabi")),
    ("Bildirim ve tebligat adresleri", ("tebligat", "bildirim adres", "adres değişikliği")),
    ("Damga vergisi ve masraflar", ("damga vergisi", "masraf", "harç")),
]
_DESEN = [(ad, re.compile(r"\b(?:" + "|".join(re.escape(k) for k in kelimeler) + ")")) for ad, kelimeler in KONULAR]

RISKLER = ["dusuk", "orta", "yuksek"]
LEHINE = ["bizim", "karsi", "dengeli", "belirsiz"]
DURUMLAR = ["var", "eksik", "belirsiz"]


def kucuk(s: str) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower()


# ----------------------------------------------------------------------------
# Maddelere bölme (deterministik)
# ----------------------------------------------------------------------------

@dataclass
class Madde:
    no: str
    baslik: str
    metin: str
    etiketler: list[str] = field(default_factory=list)


_MADDE_BASI = re.compile(r"^\s*(?:MADDE|Madde|madde)\s*[-–:.]?\s*(\d{1,3})\s*[-–:.)]?\s*(.*)$")
_NUMARALI_BASLIK = re.compile(r"^\s*(\d{1,2})[.)]\s+([A-ZÇĞİÖŞÜ][A-ZÇĞİÖŞÜ0-9 ,/&()'’\-–]{2,80})\s*$")


def maddelere_bol(metin: str) -> tuple[str, list[Madde]]:
    """'MADDE 5 – FESİH' ya da '5. FESİH' başlıklarından böler. Başlık yoksa ~2.000 karakterlik bölümler kullanır.
    Dönüş: (başlık öncesi giriş metni, maddeler)"""
    satirlar = metin.splitlines()
    for desen in (_MADDE_BASI, _NUMARALI_BASLIK):
        bas = [(i, desen.match(s)) for i, s in enumerate(satirlar)]
        bas = [(i, m) for i, m in bas if m]
        if len(bas) >= 3:
            giris = "\n".join(satirlar[:bas[0][0]]).strip()
            maddeler, gorulen = [], Counter()
            for j, (i, m) in enumerate(bas):
                son = bas[j + 1][0] if j + 1 < len(bas) else len(satirlar)
                no = m.group(1)
                gorulen[no] += 1
                if gorulen[no] > 1:                     # ekler/alt sözleşmelerde numara tekrar edebilir
                    no = f"{no}-{gorulen[no]}"
                govde = "\n".join(satirlar[i + 1:son]).strip()
                maddeler.append(Madde(no, m.group(2).strip(" -–:"), govde))
            return giris, maddeler
    # Başlık bulunamadı: paragraf gruplarına böl
    paragraflar = [p.strip() for p in re.split(r"\n\s*\n", metin) if p.strip()]
    maddeler, tampon = [], []
    for p in paragraflar:
        tampon.append(p)
        if sum(len(x) for x in tampon) > 2000:
            maddeler.append(Madde(f"B{len(maddeler) + 1}", "", "\n\n".join(tampon)))
            tampon = []
    if tampon:
        maddeler.append(Madde(f"B{len(maddeler) + 1}", "", "\n\n".join(tampon)))
    return "", maddeler


def etiketle(m: Madde) -> list[str]:
    metin = kucuk(f"{m.baslik} {m.metin}")
    return [ad for ad, desen in _DESEN if desen.search(metin)]


def gizle(metin: str, terimler: list[str]) -> str:
    metin = llm.maskele(metin)
    for i, t in enumerate(sorted({t.strip() for t in terimler if t.strip()}, key=len, reverse=True), 1):
        metin = re.sub(re.escape(t), f"[GİZLİ-{i}]", metin, flags=re.I)
    return metin


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

SEMA = {
    "type": "object",
    "properties": {
        "sozlesme_turu": {"type": "string"},
        "maddeler": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "madde": {"type": "string"},
                    "konu": {"type": "string"},
                    "ozet": {"type": "string"},
                    "risk": {"type": "string", "enum": RISKLER},
                    "lehine": {"type": "string", "enum": LEHINE},
                    "gerekce": {"type": "string"},
                    "oneri": {"type": "string"},
                    "onerilen_metin": {"type": "string"},
                },
                "required": ["madde", "konu", "ozet", "risk", "lehine", "gerekce", "oneri", "onerilen_metin"],
                "additionalProperties": False,
            },
        },
        "kontrol_listesi": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "konu": {"type": "string", "enum": [k for k, _ in KONULAR]},
                    "durum": {"type": "string", "enum": DURUMLAR},
                    "madde": {"type": "string"},
                    "not": {"type": "string"},
                },
                "required": ["konu", "durum", "madde", "not"],
                "additionalProperties": False,
            },
        },
        "oncelikli_aksiyonlar": {"type": "array", "items": {"type": "string"}},
        "genel_degerlendirme": {"type": "string"},
    },
    "required": ["sozlesme_turu", "maddeler", "kontrol_listesi", "oncelikli_aksiyonlar", "genel_degerlendirme"],
    "additionalProperties": False,
}


def paketle(maddeler: list[Madde], terimler: list[str]) -> list[list[tuple[Madde, str]]]:
    paketler, simdiki, boy = [], [], 0
    for m in maddeler:
        metin = gizle(f"{m.baslik}\n{m.metin}".strip(), terimler)
        if simdiki and boy + len(metin) > PAKET_KARAKTER:
            paketler.append(simdiki)
            simdiki, boy = [], 0
        simdiki.append((m, metin))
        boy += len(metin)
    if simdiki:
        paketler.append(simdiki)
    return paketler


def incele(giris: str, maddeler: list[Madde], taraf: str, terimler: list[str]) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    paketler = paketle(maddeler, terimler)
    sonuc = {"sozlesme_turu": "", "maddeler": {}, "kontrol": {}, "aksiyonlar": [], "yorumlar": []}
    oncelik = {"var": 2, "belirsiz": 1, "eksik": 0}
    for i, paket in enumerate(paketler, 1):
        parcali = len(paketler) > 1
        mesaj = "\n".join([
            f"<bizim_taraf>{taraf}</bizim_taraf>",
            f"<giris>{gizle(giris, terimler)}</giris>" if giris else "",
            f"<bilgi>Sözleşmenin {i}/{len(paketler)}. bölümü. Kontrol listesini yalnız bu bölüme göre doldur.</bilgi>" if parcali else "",
            "<sozlesme>",
            *[f'<madde no="{m.no}" etiketler="{", ".join(m.etiketler) or "-"}">\n{metin}\n</madde>' for m, metin in paket],
            "</sozlesme>",
            "<kontrol_listesi>" + "; ".join(k for k, _ in KONULAR) + "</kontrol_listesi>",
        ])
        yanit = llm.json_iste(sistem, mesaj, SEMA)
        gecerli = {m.no for m, _ in paket}
        for s in yanit.get("maddeler", []):
            if s.get("madde") in gecerli:               # modelin uydurduğu madde numaralarını yok say
                sonuc["maddeler"][s["madde"]] = s
        for k in yanit.get("kontrol_listesi", []):
            onceki = sonuc["kontrol"].get(k["konu"])
            if onceki is None or oncelik[k["durum"]] > oncelik[onceki["durum"]]:
                sonuc["kontrol"][k["konu"]] = k
        sonuc["aksiyonlar"] += yanit.get("oncelikli_aksiyonlar", [])
        if yanit.get("genel_degerlendirme"):
            sonuc["yorumlar"].append(yanit["genel_degerlendirme"])
        sonuc["sozlesme_turu"] = sonuc["sozlesme_turu"] or yanit.get("sozlesme_turu", "")
    return sonuc


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

RISK_AD = {"dusuk": "Düşük", "orta": "Orta", "yuksek": "Yüksek"}
LEHINE_AD = {"bizim": "Bizim lehimize", "karsi": "Karşı taraf lehine", "dengeli": "Dengeli", "belirsiz": "Belirsiz"}
DURUM_AD = {"var": "Var", "eksik": "EKSİK", "belirsiz": "Belirsiz"}
RISK_DOLGU = {"yuksek": PatternFill("solid", fgColor="FDE2E1"), "orta": PatternFill("solid", fgColor="FFF4CE"),
              "dusuk": PatternFill("solid", fgColor="E3F5E1")}
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
ONAY_DOLGU = PatternFill("solid", fgColor="FFF4CE")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
UST = Alignment(vertical="top", wrap_text=True)


def _tablo(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def rapor_yaz(cikti: Path, kaynak: Path, taraf: str, maddeler: list[Madde], sonuc: dict) -> dict:
    sirali = sorted(maddeler, key=lambda m: -RISKLER.index(sonuc["maddeler"].get(m.no, {}).get("risk", "dusuk")))
    sayac = Counter(s["risk"] for s in sonuc["maddeler"].values())
    eksikler = [k for k, _ in KONULAR if sonuc["kontrol"].get(k, {}).get("durum") == "eksik"]

    wb = Workbook()
    o = wb.active
    o.title = "Özet"
    for s in [["Sözleşme ön inceleme raporu"], [],
              ["Dosya", kaynak.name], ["İnceleme kimin adına", taraf], ["Sözleşme türü", sonuc["sozlesme_turu"]],
              ["Madde sayısı", len(maddeler)],
              ["Risk dağılımı", f"Yüksek {sayac['yuksek']} · Orta {sayac['orta']} · Düşük {sayac['dusuk']}"],
              ["Eksik düzenlemeler", ", ".join(eksikler) or "—"], []]:
        o.append(s)
    o["A1"].font = Font(bold=True, size=13)
    o.append(["Öncelikli aksiyonlar"])
    o.cell(o.max_row, 1).font = Font(bold=True)
    for i, a in enumerate(dict.fromkeys(sonuc["aksiyonlar"]), 1):
        o.append([f"{i}.", a])
    o.append([])
    for y in sonuc["yorumlar"]:
        o.append(["Genel değerlendirme", y])
    o.append(["Model", llm.kullanim_ozeti()])
    o.append(["Oluşturulma", datetime.now().strftime("%d.%m.%Y %H:%M")])
    o.append(["Önemli", "Bu rapor yapay zekâ destekli bir ön incelemedir, hukuki görüş değildir. Maddeleme ve konu "
                        "etiketleri koddan; risk, gerekçe ve öneriler modelden gelir. Avukat onayı olmadan kullanmayın."])
    o.column_dimensions["A"].width = 24
    o.column_dimensions["B"].width = 120
    for satir in o.iter_rows():
        for h in satir:
            h.alignment = UST

    r = wb.create_sheet("Madde Analizi")
    _tablo(r, ["Madde", "Başlık", "Konu", "Risk", "Kimin Lehine", "Özet", "Gerekçe", "Öneri", "Önerilen Metin", "Avukat Onayı"],
           (8, 22, 20, 9, 16, 40, 48, 44, 60, 14))
    for m in sirali:
        s = sonuc["maddeler"].get(m.no)
        if s:
            r.append([m.no, m.baslik, s["konu"], RISK_AD[s["risk"]], LEHINE_AD[s["lehine"]], s["ozet"], s["gerekce"], s["oneri"],
                      s["onerilen_metin"], ""])
            r.cell(r.max_row, 4).fill = RISK_DOLGU[s["risk"]]
            for c in range(5, 10):
                r.cell(r.max_row, c).fill = AI_DOLGU
        else:
            r.append([m.no, m.baslik, ", ".join(m.etiketler), "", "", "AI yanıt vermedi", "", "", "", ""])
        r.cell(r.max_row, 10).fill = ONAY_DOLGU
        for h in r[r.max_row]:
            h.alignment = UST
    r.auto_filter.ref = r.dimensions

    k = wb.create_sheet("Kontrol Listesi")
    _tablo(k, ["Konu", "Durum", "İlgili Madde", "Kod Etiketi (madde)", "Not", "Avukat Onayı"], (40, 11, 12, 18, 80, 14))
    for ad, _ in KONULAR:
        s = sonuc["kontrol"].get(ad, {"durum": "belirsiz", "madde": "", "not": "Model bu konuyu değerlendirmedi"})
        kod_bulgu = ", ".join(m.no for m in maddeler if ad in m.etiketler)
        k.append([ad, DURUM_AD[s["durum"]], s["madde"], kod_bulgu or "—", s["not"], ""])
        k.cell(k.max_row, 2).fill = {"eksik": RISK_DOLGU["yuksek"], "belirsiz": RISK_DOLGU["orta"], "var": RISK_DOLGU["dusuk"]}[s["durum"]]
        k.cell(k.max_row, 6).fill = ONAY_DOLGU
        for h in k[k.max_row]:
            h.alignment = UST

    t = wb.create_sheet("Sözleşme Metni")
    _tablo(t, ["Madde", "Başlık", "Metin", "Konu Etiketleri (kod)"], (8, 26, 120, 40))
    for m in maddeler:
        t.append([m.no, m.baslik, m.metin, ", ".join(m.etiketler)])
        for h in t[t.max_row]:
            h.alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

    # Kısa Markdown özeti (e-posta / not için)
    md = [f"# Sözleşme ön inceleme: {kaynak.name}", "", f"- İnceleme kimin adına: **{taraf}**",
          f"- Sözleşme türü: {sonuc['sozlesme_turu']}",
          f"- Risk dağılımı: yüksek {sayac['yuksek']}, orta {sayac['orta']}, düşük {sayac['dusuk']}",
          f"- Eksik düzenlemeler: {', '.join(eksikler) or '—'}", "", "## Öncelikli aksiyonlar", ""]
    md += [f"{i}. {a}" for i, a in enumerate(dict.fromkeys(sonuc["aksiyonlar"]), 1)]
    md += ["", "## Yüksek riskli maddeler", ""]
    for m in sirali:
        s = sonuc["maddeler"].get(m.no)
        if s and s["risk"] == "yuksek":
            md += [f"**Madde {m.no} – {m.baslik or s['konu']}** ({LEHINE_AD[s['lehine']]})", "", f"- {s['gerekce']}", f"- Öneri: {s['oneri']}", ""]
    md += ["---", "_Yapay zekâ destekli ön incelemedir; hukuki görüş değildir. Avukat değerlendirmesi gerekir._"]
    cikti.with_suffix(".md").write_text("\n".join(md), encoding="utf-8")
    return {"sayac": sayac, "eksikler": eksikler}


def calistir(girdi: Path, cikti: Path, taraf: str, terimler: list[str] | None = None, evet: bool = False) -> dict:
    metin = belge.metin_oku(girdi)
    giris, maddeler = maddelere_bol(metin)
    if not maddeler:
        raise llm.LLMHatasi(f"{girdi.name}: metin boş.")
    for m in maddeler:
        m.etiketler = etiketle(m)
    print(f"[OK] {len(maddeler)} madde/bölüm bulundu · kodla tespit edilen konu: "
          f"{len({e for m in maddeler for e in m.etiketler})}/{len(KONULAR)}")
    karakter = sum(len(m.metin) for m in maddeler)
    llm.onay_al(f"Sözleşme metni (~{karakter:,} karakter; IBAN/TCKN/telefon/e-posta".replace(",", ".")
                + (f" ve {len(terimler)} gizli terim" if terimler else "") + " maskeli) inceleme için gönderilecek.", evet)
    sonuc = incele(giris, maddeler, taraf, terimler or [])
    ozet = rapor_yaz(cikti, girdi, taraf, maddeler, sonuc)
    return {"maddeler": maddeler, "sonuc": sonuc, **ozet}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Sözleşmenin riskli maddelerini ve eksik düzenlemelerini işaretler (ön inceleme).")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "hizmet_sozlesmesi.txt", help="Sözleşme (.docx, .pdf, .txt)")
    p.add_argument("--taraf", default="", help="İncelemeyi kimin adına yapıyorsunuz? (ör. 'Müşteri', 'Alıcı', 'Kiracı')")
    p.add_argument("--gizle", nargs="*", default=[], metavar="TERİM", help="Modele gönderilmeden maskelenecek adlar/terimler")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "sozlesme_inceleme.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    if a.girdi == BURASI / "ornek_veri" / "hizmet_sozlesmesi.txt":
        a.taraf = a.taraf or "MÜŞTERİ (Hayali Lojistik)"
        a.gizle = a.gizle or ["Ayşe Örnek"]
    if not a.taraf:
        print("[X] --taraf verin: riskler sizin bakış açınızdan değerlendirilir (ör. --taraf \"Müşteri\").")
        return 2
    try:
        s = calistir(a.girdi, a.cikti, a.taraf, a.gizle, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] Risk: yüksek {s['sayac']['yuksek']} · orta {s['sayac']['orta']} · düşük {s['sayac']['dusuk']}")
    if s["eksikler"]:
        print(f"[!] Eksik düzenlemeler: {', '.join(s['eksikler'])}")
    print(f"[OK] Rapor: {a.cikti.resolve()} (+ {a.cikti.with_suffix('.md').name})")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
