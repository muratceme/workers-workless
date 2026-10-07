"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import re
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import belge  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri" / "hizmet_sozlesmesi.txt"
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    nolar = re.findall(r'<madde no="([^"]+)"', kullanici)
    return {
        "sozlesme_turu": "Hizmet sözleşmesi",
        "maddeler": [{"madde": n, "konu": "Test", "ozet": "özet", "risk": "yuksek" if n == "6" else "dusuk",
                      "lehine": "karsi" if n == "6" else "dengeli", "gerekce": "gerekçe", "oneri": "öneri",
                      "onerilen_metin": ""} for n in nolar]
        + [{"madde": "99", "konu": "Uydurma", "ozet": "", "risk": "yuksek", "lehine": "karsi", "gerekce": "",
            "oneri": "", "onerilen_metin": ""}],
        "kontrol_listesi": [{"konu": "Mücbir sebep", "durum": "eksik", "madde": "", "not": "Yok"},
                            {"konu": "Fesih", "durum": "belirsiz", "madde": "8", "not": "Tek taraflı"}],
        "oncelikli_aksiyonlar": ["Sorumluluk sınırını karşılıklı yapın."],
        "genel_degerlendirme": "Karşı taraf lehine dengesiz.",
    }


class BolmeTesti(unittest.TestCase):
    def test_madde_bolme(self):
        giris, maddeler = agent.maddelere_bol(belge.metin_oku(ORNEK))
        self.assertEqual([m.no for m in maddeler], [str(i) for i in range(1, 13)])
        self.assertEqual(maddeler[7].baslik, "FESİH")
        self.assertIn("SÖZLEŞMESİ", giris)

    def test_numarali_baslik_ve_tekrar(self):
        metin = "1. TARAFLAR\nA ile B\n2. KONU\nmal alımı\n3. FESİH\nfesih\nEK-1\n1. TANIMLAR\nx"
        _, maddeler = agent.maddelere_bol(metin)
        self.assertEqual([m.no for m in maddeler], ["1", "2", "3", "1-2"])

    def test_basliksiz_metin(self):
        _, maddeler = agent.maddelere_bol("Paragraf bir.\n\nParagraf iki.")
        self.assertEqual([m.no for m in maddeler], ["B1"])

    def test_etiket(self):
        m = agent.Madde("7", "GİZLİLİK", "Bilgileri gizli tutar. Aksi hâlde cezai şart öder.")
        self.assertIn("Gizlilik", agent.etiketle(m))
        self.assertIn("Ceza şartı", agent.etiketle(m))

    def test_docx(self):
        with tempfile.TemporaryDirectory() as tmp:
            yol = Path(tmp) / "s.docx"
            ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            xml = (f'<w:document xmlns:w="{ns}"><w:body><w:p><w:r><w:t>MADDE 1 – TARAFLAR</w:t></w:r></w:p>'
                   f'<w:p><w:r><w:t>A ve B</w:t></w:r></w:p></w:body></w:document>')
            with zipfile.ZipFile(yol, "w") as z:
                z.writestr("word/document.xml", xml)
            self.assertEqual(belge.metin_oku(yol), "MADDE 1 – TARAFLAR\nA ve B")
            with self.assertRaises(belge.BelgeHatasi):
                belge.metin_oku(Path(tmp) / "eski.doc")


class AgentTesti(unittest.TestCase):
    def setUp(self):
        MESAJLAR.clear()

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_uctan_uca(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "s.xlsx"
            s = agent.calistir(ORNEK, cikti, "Müşteri", ["Ayşe Örnek"], evet=True)
            self.assertEqual(len(MESAJLAR), 1)
            m = MESAJLAR[0]
            self.assertNotIn("Ayşe Örnek", m)                     # --gizle
            self.assertNotIn("ayse.ornek@", m)                    # e-posta maskesi
            self.assertNotIn("0532 000 00 00", m)
            self.assertIn("[IBAN]", m)
            self.assertIn("<bizim_taraf>Müşteri</bizim_taraf>", m)
            self.assertNotIn("99", s["sonuc"]["maddeler"])        # uydurma madde yok sayılır
            self.assertEqual(s["sayac"]["yuksek"], 1)
            self.assertEqual(s["eksikler"], ["Mücbir sebep"])

            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Özet", "Madde Analizi", "Kontrol Listesi", "Sözleşme Metni"])
            ws = wb["Madde Analizi"]
            self.assertEqual(ws.cell(2, 1).value, "6")              # yüksek riskli madde en üstte
            self.assertEqual(ws.cell(1, ws.max_column).value, "Avukat Onayı")
            self.assertEqual(wb["Kontrol Listesi"].max_row, len(agent.KONULAR) + 1)
            md = cikti.with_suffix(".md").read_text(encoding="utf-8")
            self.assertIn("Madde 6", md)

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_uzun_sozlesme_paketlenir(self, _):
        with mock.patch.object(agent, "PAKET_KARAKTER", 1500), tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(ORNEK, Path(tmp) / "s.xlsx", "Müşteri", evet=True)
        self.assertGreater(len(MESAJLAR), 1)
        self.assertEqual(len(s["sonuc"]["maddeler"]), 12)          # her madde bir kez
        self.assertIn("bölümü", MESAJLAR[0])

    def test_sema_konulari(self):
        konular = agent.SEMA["properties"]["kontrol_listesi"]["items"]["properties"]["konu"]["enum"]
        self.assertEqual(len(konular), 20)


if __name__ == "__main__":
    unittest.main()
