"""
Taşıma Talebi Bilgi Çıkarma — Workers / Workless AI Agent
Lojistik › Havayolu Operasyon Uzmanı (Depo / Lojistik Uzmanı da kullanır)

Müşterilerden ya da iç birimlerden e-postayla gelen taşıma taleplerini planlama ve teklif tablolarına aktarır:
uluslararası navlun teklif talepleri (hava / deniz) ve yurt içi sevkiyat talepleri.
1. Kod talep belgelerini okur; telefon, e-posta, IBAN ve kimlik numaraları maskelenir.
2. Model talebin türünü, müşteri / referansı, çıkış ve varışı, teslim şeklini, taşıma modunu, hazır / teslim tarihini,
   mal tanımını, tehlikeli madde ve istif bilgisini, kapları (adet, ölçü, ağırlık; ağırlık kap başına mı toplam mı) ve
   yurt içi sevkiyat satırlarını belgedeki yazımla çıkarır.
3. Kod çıkarımı denetler: her sayı ve tarih belgede geçiyor mu. Sorun varsa modele geri gönderilir (en fazla --tur kez).
4. Kod hesaplar: toplam ağırlık verilmişse kap başına ağırlık (toplam ÷ adet), toplam brüt ağırlık ve hacim (m³), hava için
   hacimsel ağırlık (÷ --bolen, varsayılan 6000); belgede yazan toplamla tutarsızlık, eksik ölçü / ağırlık, tehlikeli madde
   (DGR / IMDG beyanı gerekir) ve istiflenemez yük bulguları.
Çıktı: her navlun talebi için "Navlun Teklif Hesaplama" kod bloğunun kap dosyası (<ref>_kaplar.xlsx) ve yurt içi talepler
için "Sevkiyat Araç Planlama" kod bloğunun sipariş dosyası (sevkiyat_siparisleri.xlsx).

Kullanım:
    python agent.py                                                    # örnek: hava, deniz ve yurt içi talep
    python agent.py --girdi ./talepler --bolen 6000
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
RENK = {"Hata": "FDE2E1", "Dikkat": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1 or (s.count(".") == 1 and len(s.split(".")[1]) == 3):
        s = s.replace(".", "")
    try:
        return Decimal(s) if s not in ("", "-", ".") else None
    except InvalidOperation:
        return None


def sayilar(metin: str) -> set[Decimal]:
    temiz = re.sub(r"\b\d{1,2}\.\d{1,2}\.\d{4}\b", " ", str(metin))
    temiz = re.sub(r"(\d)\s*[x×]\s*(?=\d)", r"\1 ", temiz)                      # 60x40x40 → 60 40 40
    return {v for x in re.findall(r"\d+(?:[.,]\d+)*", temiz) if (v := para(x)) is not None}


def tarih_coz(x) -> date | None:
    m = re.search(r"\b(\d{1,2})\.(\d{1,2})\.(\d{4})\b", str(x or ""))
    if not m:
        return None
    try:
        return date(int(m[3]), int(m[2]), int(m[1]))
    except ValueError:
        return None


def tr(x: Decimal, ondalik: int = 2) -> str:
    return f"{x:,.{ondalik}f}".replace(",", "X").replace(".", ",").replace("X", ".")


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

KAP = {"type": "object", "properties": {"kap": S, "adet": S, "boy": S, "en": S, "yukseklik": S, "agirlik": S,
                                        "agirlik_tipi": {"type": "string", "enum": ["Kap başına", "Toplam", "Belirtilmemiş"]}},
       "required": ["kap", "adet", "boy", "en", "yukseklik", "agirlik", "agirlik_tipi"], "additionalProperties": False}
SEVK = {"type": "object", "properties": {"siparis_no": S, "musteri": S, "bolge": S, "teslim_tarihi": S, "palet": S, "agirlik": S, "hacim": S},
        "required": ["siparis_no", "musteri", "bolge", "teslim_tarihi", "palet", "agirlik", "hacim"], "additionalProperties": False}
SEMA = {"type": "object", "properties": {
    "tur": {"type": "string", "enum": ["Navlun teklifi", "Yurt içi sevkiyat", "Taşıma talebi değil"]}, "musteri": S, "referans": S,
    "cikis": S, "varis": S, "teslim_sekli": S, "mod": {"type": "string", "enum": ["Hava", "Deniz", "Karayolu", "Belirtilmemiş"]},
    "hazir_tarihi": S, "mal_tanimi": S, "tehlikeli_madde": S, "istiflenemez": {"type": "boolean"}, "belgedeki_toplam_agirlik": S,
    "kaplar": {"type": "array", "items": KAP}, "sevkiyatlar": {"type": "array", "items": SEVK}, "notlar": {"type": "array", "items": S}},
    "required": ["tur", "musteri", "referans", "cikis", "varis", "teslim_sekli", "mod", "hazir_tarihi", "mal_tanimi", "tehlikeli_madde", "istiflenemez",
                 "belgedeki_toplam_agirlik", "kaplar", "sevkiyatlar", "notlar"], "additionalProperties": False}


def denetle(c: dict, metin: str) -> list[str]:
    if c["tur"] == "Taşıma talebi değil":
        return []
    sorun, say = [], sayilar(metin)

    def sayi(on, ad, v):
        if v.strip() and (para(v) is None or para(v) not in say):
            sorun.append(f"{on}: {ad} '{v}' belgede yazmıyor (hesaplamayın; yoksa boş bırakın)")

    def tarih(on, ad, v):
        if v.strip():
            t = tarih_coz(v)
            if t is None or f"{t:%d.%m.%Y}" not in metin:
                sorun.append(f"{on}: {ad} '{v}' belgede GG.AA.YYYY olarak yazmıyor")

    for a, ad in (("musteri", "Müşteri"), ("referans", "Referans"), ("cikis", "Çıkış"), ("varis", "Varış"), ("teslim_sekli", "Teslim şekli")):
        if c[a].strip() and not metinde_var(c[a], metin):
            sorun.append(f"Talep: {ad} '{c[a][:40]}' belgede geçmiyor")
    tarih("Talep", "Hazır tarihi", c["hazir_tarihi"])
    sayi("Talep", "Toplam ağırlık", c["belgedeki_toplam_agirlik"])
    for i, k in enumerate(c["kaplar"], 1):
        for a in ("adet", "boy", "en", "yukseklik", "agirlik"):
            sayi(f"Kap {i}", a, k[a])
        if para(k["adet"]) is None:
            sorun.append(f"Kap {i}: adet boş")
    for i, s in enumerate(c["sevkiyatlar"], 1):
        if s["siparis_no"].strip() and not metinde_var(s["siparis_no"], metin):
            sorun.append(f"Sevkiyat {i}: sipariş no '{s['siparis_no']}' belgede geçmiyor")
        for a in ("palet", "agirlik", "hacim"):
            sayi(f"Sevkiyat {i}", a, s[a])
        tarih(f"Sevkiyat {i}", "Teslim tarihi", s["teslim_tarihi"])
    if c["tur"] == "Navlun teklifi" and not c["kaplar"] and not c["notlar"]:
        sorun.append("Talep: navlun teklifinde kap satırı yok; belgede kap bilgisi yoksa notlara yazın")
    if c["tur"] == "Yurt içi sevkiyat" and not c["sevkiyatlar"]:
        sorun.append("Talep: yurt içi sevkiyatta sipariş satırı yok")
    return sorun


def hesap(c: dict, bolen: Decimal) -> dict:
    """Kod: kap başına ağırlık, toplam brüt, hacim, hacimsel ağırlık ve bulgular."""
    b, kaplar, brut, cm3 = [], [], Decimal(0), Decimal(0)
    for i, k in enumerate(c["kaplar"], 1):
        adet = para(k["adet"]) or Decimal(0)
        olc = [para(k[a]) for a in ("boy", "en", "yukseklik")]
        kg = para(k["agirlik"])
        if kg is not None and k["agirlik_tipi"] == "Toplam" and adet:
            birim = (kg / adet).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            b.append(("Bilgi", f"Kap {i} ({k['kap']}): toplam {tr(kg)} kg ÷ {adet:g} = kap başına {tr(birim)} kg"))
            kg = birim
        elif kg is not None and k["agirlik_tipi"] == "Belirtilmemiş":
            b.append(("Dikkat", f"Kap {i} ({k['kap']}): ağırlığın kap başına mı toplam mı olduğu belirsiz; kap başına sayıldı"))
        if None in olc:
            b.append(("Dikkat", f"Kap {i} ({k['kap']}): ölçü eksik"))
        if kg is None:
            b.append(("Dikkat", f"Kap {i} ({k['kap']}): ağırlık eksik"))
        if None not in olc:
            cm3 += adet * olc[0] * olc[1] * olc[2]
        if kg is not None:
            brut += adet * kg
        kaplar.append({"kap": k["kap"] or f"Kap {i}", "adet": adet, "boy": olc[0], "en": olc[1], "yuk": olc[2], "kg": kg})
    bt = para(c["belgedeki_toplam_agirlik"])
    if bt is not None and kaplar and abs(bt - brut) > Decimal("0.5"):
        b.append(("Hata", f"Belgedeki toplam brüt {tr(bt)} kg, kaplardan hesaplanan {tr(brut)} kg: müşteriye teyit ettirin"))
    if c["tehlikeli_madde"].strip():
        kural = "IATA DGR" if c["mod"] == "Hava" else "IMDG" if c["mod"] == "Deniz" else "ADR"
        b.append(("Dikkat", f"Tehlikeli madde: {c['tehlikeli_madde']} — {kural} kapsamında beyan, ambalaj ve taşıyıcı onayı gerekir"))
    if c["istiflenemez"]:
        b.append(("Dikkat", "İstiflenemez yük: taşıyıcı istiflenemezlik ek ücreti / yer hesabı uygulayabilir"))
    for i, s in enumerate(c["sevkiyatlar"], 1):
        eksik = [ad for a, ad in (("palet", "palet"), ("agirlik", "ağırlık"), ("hacim", "hacim"), ("teslim_tarihi", "teslim tarihi")) if not s[a].strip()]
        if eksik:
            b.append(("Dikkat", f"Sevkiyat {s['siparis_no'] or i}: {', '.join(eksik)} eksik; araç planına eksik bilgiyle girer"))
    m3 = (cm3 / 1000000).quantize(Decimal("0.001"), rounding=ROUND_HALF_UP)
    hacimsel = (cm3 / bolen).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP) if cm3 else Decimal(0)
    return {"kaplar": kaplar, "brut": brut, "m3": m3, "hacimsel": hacimsel, "bulgu": b}


def uret(metin: str, duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    mesaj = f"<talep>\n{metin}\n</talep>"
    if duzeltme:
        mesaj += ("\n<duzeltme>\nÖnceki çıkarım: " + json.dumps(onceki, ensure_ascii=False) + "\nKodun bulduğu sorunlar:\n"
                  + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")
    return llm.json_iste(sistem, mesaj, SEMA)


def duzeltme_dongusu(uret_f, denetle_f, tur: int) -> tuple[dict, list[str], int]:
    taslak = uret_f(None, None)
    kalan = denetle_f(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret_f(kalan, taslak)
        kalan = denetle_f(taslak)
    return taslak, kalan, yapilan


# ----------------------------------------------------------------------------
# Çıktı
# ----------------------------------------------------------------------------

def _baslik(ws, basliklar, genislik=()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font, h.alignment = BASLIK_DOLGU, BASLIK_YAZI, UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def _f(x):
    return None if x is None else float(x)


def dosya_adi(s: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", katla(s).replace(" ", "-"))[:40] or "talep"


def rapor_yaz(cikti: Path, talepler: list[dict], bolen: Decimal) -> list[Path]:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Talepler"
    _baslik(ws, ["Kaynak Dosya", "Tür", "Mod", "Müşteri", "Referans", "Çıkış", "Varış", "Teslim Şekli", "Hazır Tarihi", "Mal", "Kap Adedi", "Brüt (kg)",
                 "Hacim (m³)", f"Hacimsel Ağırlık (kg, ÷{bolen:g})", "Bulgular", "Model Notları", "Kap Dosyası", "Son Kontrol"],
            (26, 16, 9, 20, 14, 18, 18, 14, 11, 28, 8, 10, 10, 12, 60, 40, 26, 30))
    ek = []
    sevk = []
    for t in talepler:
        c, h = t["c"], t["h"]
        kap_dosya = ""
        if c["tur"] == "Navlun teklifi" and h["kaplar"]:
            kw = Workbook()
            ks = kw.active
            ks.title = "Kaplar"
            _baslik(ks, ["Kap", "Adet", "Boy (cm)", "En (cm)", "Yükseklik (cm)", "Ağırlık (kg)"], (24, 8, 9, 9, 13, 12))
            for k in h["kaplar"]:
                ks.append([k["kap"], _f(k["adet"]), _f(k["boy"]), _f(k["en"]), _f(k["yuk"]), _f(k["kg"])])
            y = cikti.with_name(f"{dosya_adi(c['referans'] or Path(t['dosya']).stem)}_kaplar.xlsx")
            kw.save(y)
            ek.append(y)
            kap_dosya = y.name
        for s in c["sevkiyatlar"]:
            sevk.append([s["siparis_no"], s["musteri"], s["bolge"], tarih_coz(s["teslim_tarihi"]), _f(para(s["palet"])), _f(para(s["agirlik"])),
                         _f(para(s["hacim"]))])
        ws.append([t["dosya"], c["tur"], c["mod"], c["musteri"], c["referans"], c["cikis"], c["varis"], c["teslim_sekli"], tarih_coz(c["hazir_tarihi"]),
                   c["mal_tanimi"], float(sum((k["adet"] for k in h["kaplar"]), Decimal(0))) if h["kaplar"] else None, _f(h["brut"]) if h["kaplar"] else None,
                   _f(h["m3"]) if h["kaplar"] else None, _f(h["hacimsel"]) if h["kaplar"] else None, "\n".join(f"{o}: {m}" for o, m in h["bulgu"]),
                   "\n".join(c["notlar"]), kap_dosya, "; ".join(t["kalan"]) or "Geçti"])
        n = ws.max_row
        ws.cell(n, 9).number_format = "DD.MM.YYYY"
        for j in (2, 3, 4, 5, 6, 7, 8, 10):
            ws.cell(n, j).fill = AI_DOLGU
        if any(o == "Hata" for o, _ in h["bulgu"]):
            ws.cell(n, 15).fill = KIRMIZI
        ws.cell(n, 18).fill = KIRMIZI if t["kalan"] else YESIL
        for x in ws[n]:
            x.alignment = UST
    s = wb.create_sheet("Sevkiyat Siparişleri")
    _baslik(s, ["Sipariş No", "Müşteri", "Teslim Bölgesi", "Teslim Tarihi", "Palet", "Ağırlık (kg)", "Hacim (m3)"], (12, 26, 16, 12, 7, 11, 10))
    for r in sevk:
        s.append(r)
        s.cell(s.max_row, 4).number_format = "DD.MM.YYYY"
    k = wb.create_sheet("Bilgi")
    k.append([f"Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    k.append(["Önemli: Mor hücreler yapay zekâ ile çıkarılmıştır; teklif vermeden önce müşterinin talebiyle karşılaştırın. Toplam ağırlık, hacim ve "
              "hacimsel ağırlık kodla hesaplanmıştır. Kap dosyaları 'Navlun Teklif Hesaplama' (--kaplar), sevkiyat_siparisleri.xlsx 'Sevkiyat Araç "
              "Planlama' (--siparisler) kod bloğuna verilebilir."])
    wb.save(cikti)
    if sevk:
        sw = Workbook()
        ss = sw.active
        _baslik(ss, ["Sipariş No", "Müşteri", "Teslim Bölgesi", "Teslim Tarihi", "Palet", "Ağırlık (kg)", "Hacim (m3)"])
        for r in sevk:
            ss.append(r)
            ss.cell(ss.max_row, 4).number_format = "DD.MM.YYYY"
        y = cikti.with_name("sevkiyat_siparisleri.xlsx")
        sw.save(y)
        ek.append(y)
    return ek


def calistir(girdi: Path, cikti: Path, bolen: Decimal = Decimal(6000), tur: int = 2, evet: bool = False) -> dict:
    if girdi.is_dir():
        dosyalar = sorted(p for p in girdi.iterdir() if p.is_file() and p.suffix.lower() in belge.DESTEKLENEN)
    elif girdi.exists():
        dosyalar = [girdi]
    else:
        raise belge.BelgeHatasi(f"{girdi}: bulunamadı.")
    if not dosyalar:
        raise belge.BelgeHatasi(f"{girdi}: talep belgesi yok.")
    print(f"[OK] {len(dosyalar)} talep belgesi")
    llm.onay_al(f"{len(dosyalar)} taşıma talebinin metni gönderilecek (telefon, e-posta, IBAN, TCKN maskeli).", evet)
    talepler = []
    for p in dosyalar:
        metin = llm.maskele(belge.metin_oku(p))
        c, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(metin, d, o), lambda x: denetle(x, metin), tur)
        talepler.append({"dosya": p.name, "c": c, "h": hesap(c, bolen), "kalan": kalan, "tur": yapilan})
    ek = rapor_yaz(cikti, talepler, bolen)
    return {"talepler": talepler, "ek": ek}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass
    p = argparse.ArgumentParser(description="Taşıma taleplerini (navlun teklifi, yurt içi sevkiyat) teklif ve araç planlama tablolarına aktarır.")
    p.add_argument("--girdi", type=Path, default=ORNEK / "talepler", help="Talep belgesi ya da klasörü (.pdf, .docx, .txt)")
    p.add_argument("--bolen", default="6000", help="Hava kargo hacimsel ağırlık böleni (cm³/kg; varsayılan 6000)")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "tasima_talepleri.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    bolen = para(a.bolen)
    if bolen is None or bolen <= 0:
        print("[X] --bolen pozitif olmalı")
        return 2
    try:
        s = calistir(a.girdi, a.cikti, bolen, a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    for t in s["talepler"]:
        c, h = t["c"], t["h"]
        ek = f" · {len(h['kaplar'])} kap · {tr(h['brut'])} kg · {tr(h['m3'], 3)} m³" if h["kaplar"] else f" · {len(c['sevkiyatlar'])} sevkiyat"
        print(f"[OK] {t['dosya']} · {c['tur']} · {c['mod']}{ek}")
        for o, m in h["bulgu"]:
            if o != "Bilgi":
                print(f"[!]   {o}: {m}")
    print(f"[OK] Rapor: {a.cikti.resolve()}" + (f" (+ {', '.join(y.name for y in s['ek'])})" if s["ek"] else ""))
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
