"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import json
import os
import sys
import tempfile
import unittest
from datetime import datetime
from decimal import Decimal as D
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri" / "talepler"
CEVAP = json.loads((Path(__file__).parent / "sahte_cevaplar.json").read_text(encoding="utf-8"))
HAVA = llm.maskele((ORNEK / "kurgu_medikal_hava.txt").read_text(encoding="utf-8"))
MESAJLAR = []


def hava_hatali() -> dict:
    """Sık model hataları: belgedeki tutarsız toplamı 'düzeltmek', koli satırına toplam ağırlığı hesaplayıp yazmak."""
    c = copy.deepcopy(CEVAP["hava"])
    c["belgedeki_toplam_agirlik"] = "264"
    c["kaplar"][0].update(agirlik="54", agirlik_tipi="Toplam")
    return c


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    if "Frankfurt" in kullanici:
        return copy.deepcopy(CEVAP["hava"] if "<duzeltme>" in kullanici else hava_hatali())
    return copy.deepcopy(CEVAP["deniz" if "Hamburg" in kullanici else "yurtici"])


class KodTesti(unittest.TestCase):
    def test_sayilar(self):
        self.assertTrue({D(60), D(40), D(18), D(120), D(80), D(100), D(210), D(254), D(3481)} <= agent.sayilar(HAVA))
        self.assertEqual(agent.sayilar("120 x 100 x 150 cm, 3,6 m3"), {D(120), D(100), D(150), D("3.6"), D(3)})

    def test_hesap(self):
        h = agent.hesap(CEVAP["hava"], D(6000))
        self.assertEqual((h["brut"], h["m3"], h["hacimsel"]), (D(264), D("1.248"), D("208.0")))
        self.assertEqual([o for o, _ in h["bulgu"]], ["Hata", "Dikkat"])
        self.assertIn("IATA DGR", h["bulgu"][1][1])
        h = agent.hesap(CEVAP["deniz"], D(6000))
        self.assertEqual((h["kaplar"][0]["kg"], h["brut"], h["m3"]), (D(300), D(1800), D("10.800")))
        self.assertEqual([o for o, _ in h["bulgu"]], ["Bilgi", "Dikkat"])
        self.assertTrue(any("IMDG" in m for _, m in agent.hesap(CEVAP["deniz"] | {"tehlikeli_madde": "Boya, UN1263"}, D(6000))["bulgu"]))
        h = agent.hesap(CEVAP["yurtici"], D(6000))
        self.assertEqual(h["bulgu"], [("Dikkat", "Sevkiyat SP-2703: ağırlık, hacim eksik; araç planına eksik bilgiyle girer")])
        c = copy.deepcopy(CEVAP["hava"])
        c["kaplar"][1].update(boy="", agirlik_tipi="Belirtilmemiş")
        self.assertEqual([o for o, _ in agent.hesap(c, D(5000))["bulgu"]][:2], ["Dikkat", "Dikkat"])

    def test_denetim(self):
        self.assertEqual(agent.denetle(CEVAP["hava"], HAVA), [])
        s = agent.denetle(hava_hatali(), HAVA)
        self.assertEqual(s, ["Talep: Toplam ağırlık '264' belgede yazmıyor (hesaplamayın; yoksa boş bırakın)",
                             "Kap 1: agirlik '54' belgede yazmıyor (hesaplamayın; yoksa boş bırakın)"])
        y = llm.maskele((ORNEK / "satis_yurtici_sevkiyat.txt").read_text(encoding="utf-8"))
        self.assertEqual(agent.denetle(CEVAP["yurtici"], y), [])
        c = copy.deepcopy(CEVAP["yurtici"])
        c["sevkiyatlar"][0].update(teslim_tarihi="24.10.2026", siparis_no="SP-9999")
        self.assertEqual(len(agent.denetle(c, y)), 2)
        self.assertEqual(agent.denetle({"tur": "Taşıma talebi değil"}, y), [])


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ornek_ve_dosyalar(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            self.assertEqual(agent.main(["--evet", "--cikti", str(cikti)]), 0)
            self.assertEqual(len(MESAJLAR), 4)                           # 3 talep + hava için 1 düzeltme
            k = [list(r) for r in load_workbook(Path(tmp) / "km-2026-88_kaplar.xlsx").active.iter_rows(values_only=True)]
            self.assertEqual(k, [["Kap", "Adet", "Boy (cm)", "En (cm)", "Yükseklik (cm)", "Ağırlık (kg)"], ["Koli", 3, 60, 40, 40, 18], ["Palet", 1, 120, 80, 100, 210]])
            d = [list(r) for r in load_workbook(Path(tmp) / "ornek-mobilya-deniz_kaplar.xlsx").active.iter_rows(values_only=True)]
            self.assertEqual(d[1], ["Paletli sandık", 6, 120, 100, 150, 300])
            s = [list(r) for r in load_workbook(Path(tmp) / "sevkiyat_siparisleri.xlsx").active.iter_rows(values_only=True)]
            self.assertEqual(s[0], ["Sipariş No", "Müşteri", "Teslim Bölgesi", "Teslim Tarihi", "Palet", "Ağırlık (kg)", "Hacim (m3)"])
            self.assertEqual(s[1], ["SP-2701", "Deneme Market İzmir", "İzmir", datetime(2026, 10, 22), 3, 1150, 3.6])
            ws = load_workbook(cikti)["Talepler"]
            self.assertEqual([ws.cell(2, j).value for j in (2, 3, 12, 13, 14)], ["Navlun teklifi", "Hava", 264, 1.248, 208])

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(agent.main(["--bolen", "0", "--evet"]), 2)
            self.assertEqual(agent.main(["--girdi", str(Path(tmp) / "yok"), "--evet"]), 1)
            s = agent.calistir(ORNEK / "kurgu_medikal_hava.txt", Path(tmp) / "t.xlsx", tur=0, evet=True)
            self.assertEqual(len(s["talepler"][0]["kalan"]), 2)


if __name__ == "__main__":
    unittest.main()
