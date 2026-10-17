"""
Ziyaret Notu Özeti — Gerisi Bende AI Agent
Satış › Satış Temsilcisi

1. Kod not dosyasını ziyaretlere ayırır ("Ziyaret: Müşteri | GG.AA.YYYY" satırı) ve her ziyaretin satırlarını numaralar.
   Telefon, e-posta ve kimlik numaraları maskelenir.
2. Model her ziyaret için CRM'e girilecek özet, müşteri ihtiyaçları, itirazlar, rakip bilgisi, fırsat aşaması ve takip
   aksiyonlarını (sorumlu, termin) yazar; her madde dayandığı not satırlarını gösterir.
3. Kod denetler:
   - Her maddenin dayanak satırı var mı ve geçerli mi; madde ve tutardaki sayılar dayandığı satırlarda yazıyor mu
   - Aksiyon termini notta yazan tarih ya da ziyaret tarihine göre "yarın", "2 gün içinde", "haftaya cuma", "ay sonu"
     gibi ifadeden hesaplanabilen bir tarih mi (termin uydurulmasın)
   - "Rakip" geçen satırlar rakip bilgisinde kullanılmış mı
   Sorun varsa yalnız o ziyaretin özeti modele geri gönderilir (en fazla --tur kez).
4. Kod; indirim / fiyat taahhüdü içeren aksiyonları "yönetici onayı gerekir", termini olmayanları "termin belirsiz" diye işaretler.

Kullanım:
    python agent.py                                               # örnek ziyaret notları
    python agent.py --girdi notlar.txt --rakipler "X Ambalaj,Y Kutu"
"""
from __future__ import annotations

import argparse
import calendar
import csv
import json
import re
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import belge
import llm

BURASI = Path(__file__).resolve().parent
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
AI_DOLGU = PatternFill("solid", fgColor="EDE7FF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
UST = Alignment(vertical="top", wrap_text=True)
S = {"type": "string"}


def katla(s) -> str:
    """Karşılaştırma için: Türkçe harfleri sadeleştirir, küçük harfe çevirir, noktalamayı boşluk yapar."""
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def metinde_var(alinti: str, metin: str) -> bool:
    """Alıntı (boşluk, noktalama ve büyük/küçük harf farkı gözetilmeden) metinde geçiyor mu?"""
    a = katla(alinti)
    return bool(a) and a in katla(metin)


def sayilar(metin) -> set[str]:
    """Metindeki sayılar ('1.250,50' → '1250.5'; '08:30' → '08:30'; '15.10.2026' → '15.10.2026')."""
    s = str(metin or "")
    out = set(re.findall(r"\b\d{1,2}[.:]\d{2}(?:[.:]\d{2,4})?\b", s))
    for x in re.findall(r"\d+(?:[.,]\d+)*", s):
        out.add(sayi_norm(x))
    return out


def sayi_norm(x: str) -> str:
    x = str(x).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", x):
        x = x.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", x):
        x = x.replace(",", "")
    else:
        x = x.replace(",", ".")
    try:
        v = float(x)
    except ValueError:
        return x
    return sayi_yaz(v)


def sayi_yaz(v) -> str:
    """Üstel gösterim olmadan sade sayı: 19475.5 → '19475.5', 8450.0 → '8450'."""
    return f"{float(v):.6f}".rstrip("0").rstrip(".")


def tarih_coz(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def zaman_coz(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: dosya bulunamadı.")
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            s = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:5])
        s = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in s if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, alanlar: dict[str, tuple], zorunlu: tuple = ()) -> list[dict]:
    """Başlık satırını eşanlamlılarla eşler; her satırı {alan: değer} sözlüğüne çevirir (boş hücre → '')."""
    s = tablo_oku(yol)
    if not s:
        raise llm.LLMHatasi(f"{yol.name}: tablo boş.")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in alanlar.items()}
    eksik = [alanlar[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise llm.LLMHatasi(f"{yol.name}: gerekli sütun bulunamadı: {', '.join(eksik)}. Başlıklar: {s[0]}")
    out = []
    for r in s[1:]:
        k = {a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()}
        k = {a: (v.strip() if isinstance(v, str) else v) for a, v in k.items()}
        out.append(k)
    return out


def baslik(ws, basliklar: list, genislik: tuple = ()) -> None:
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hizala(ws, ilk: int = 2) -> None:
    for satir in ws.iter_rows(min_row=ilk):
        for h in satir:
            h.alignment = UST


def boya(ws, sutunlar, dolgu) -> None:
    for c in sutunlar:
        ws.cell(ws.max_row, c).fill = dolgu


def duzeltme_dongusu(uret, denetle, tur: int) -> tuple[dict, list[str], int]:
    """uret(duzeltme, onceki) → taslak; denetle(taslak) → sorunlar. Sorun kalırsa en fazla `tur` kez yeniden yazdırır."""
    taslak = uret(None, None)
    kalan = denetle(taslak)
    yapilan = 0
    while kalan and yapilan < tur:
        yapilan += 1
        print(f"[i] Düzeltme turu {yapilan}: {len(kalan)} sorun modele geri gönderiliyor")
        taslak = uret(kalan, taslak)
        kalan = denetle(taslak)
    return taslak, kalan, yapilan


def duzeltme_blogu(duzeltme, onceki) -> str:
    if not duzeltme:
        return ""
    return ("\n<duzeltme>\nÖnceki taslak: " + json.dumps(onceki, ensure_ascii=False, default=str) + "\nKodun bulduğu sorunlar:\n"
            + "\n".join(f"- {s}" for s in duzeltme) + "\n</duzeltme>")


def ortam_hazirla() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    try:
        from dotenv import load_dotenv
        load_dotenv(BURASI / ".env")
        load_dotenv()
    except ImportError:
        pass


def son_not(ws, metin: str, kalan: list[str], tur: int) -> None:
    ws.append([])
    ws.append(["Son kontrol: " + ("; ".join(kalan) if kalan else "Kod denetiminden sorunsuz geçti")])
    ws.cell(ws.max_row, 1).fill = KIRMIZI if kalan else YESIL
    ws.append([f"Düzeltme turu: {tur} · Model: {llm.kullanim_ozeti()} · {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([metin])
K = {"type": "array", "items": {"type": "integer"}}
_M = {"type": "object", "properties": {"metin": S, "satirlar": K}, "required": ["metin", "satirlar"], "additionalProperties": False}
SEMA = {
    "type": "object",
    "properties": {
        "ozet": S,
        "ihtiyaclar": {"type": "array", "items": _M},
        "itirazlar": {"type": "array", "items": _M},
        "rakip_bilgisi": {"type": "array", "items": _M},
        "firsat": {"type": "object", "properties": {
            "asama": {"type": "string", "enum": ["İlk görüşme", "İhtiyaç belirleme", "Teklif", "Pazarlık", "Kapanış", "Kaybedildi"]},
            "tahmini_tutar": S, "satirlar": K}, "required": ["asama", "tahmini_tutar", "satirlar"], "additionalProperties": False},
        "aksiyonlar": {"type": "array", "items": {"type": "object", "properties": {
            "is": S, "sorumlu": {"type": "string", "enum": ["Biz", "Müşteri"]}, "termin": S, "satirlar": K},
            "required": ["is", "sorumlu", "termin", "satirlar"], "additionalProperties": False}},
    },
    "required": ["ozet", "ihtiyaclar", "itirazlar", "rakip_bilgisi", "firsat", "aksiyonlar"],
    "additionalProperties": False,
}
BASLIK = re.compile(r"^\s*Ziyaret\s*:\s*(.+?)\s*\|\s*(\S+)\s*$", re.I)
GUNLER = ["pazartesi", "sali", "carsamba", "persembe", "cuma", "cumartesi", "pazar"]
AYLAR = ["ocak", "subat", "mart", "nisan", "mayis", "haziran", "temmuz", "agustos", "eylul", "ekim", "kasim", "aralik"]
ONAY = re.compile(r"(indirim|iskonto|%\s*\d|\d+\s*%|fiyat sabit|ücretsiz|bedelsiz|vade farkı|garanti)", re.I)


def ziyaretleri_oku(yol: Path) -> list[dict]:
    if not yol.exists():
        raise llm.LLMHatasi(f"{yol}: not dosyası bulunamadı.")
    out = []
    for s in belge.metin_oku(yol).splitlines():
        m = BASLIK.match(s)
        if m:
            t = tarih_coz(m.group(2))
            if not t:
                raise llm.LLMHatasi(f"'{s.strip()}': ziyaret tarihi okunamadı (GG.AA.YYYY).")
            out.append({"kimlik": f"V{len(out) + 1}", "musteri": m.group(1), "tarih": t, "satirlar": []})
        elif s.strip() and out:
            out[-1]["satirlar"].append(llm.maskele(s.strip()))
    if not out:
        raise llm.LLMHatasi(f"{yol.name}: 'Ziyaret: Müşteri | GG.AA.YYYY' başlıklı ziyaret yok.")
    return [z for z in out if z["satirlar"]]


def _sonraki_gun(d: date, gun: int) -> date:
    return d + timedelta(days=(gun - d.weekday()) % 7 or 7)


def gecerli_terminler(metin: str, d: date) -> set[date]:
    """Notta yazan ya da ziyaret tarihine göre göreli ifadeden hesaplanabilen tarihler."""
    k, out = katla(metin), set()
    for g, a, y in re.findall(r"\b(\d{1,2})[./](\d{1,2})(?:[./](\d{2,4}))?\b", metin):
        try:
            yil = int(y) + (2000 if y and len(y) == 2 else 0) if y else d.year
            t = date(yil, int(a), int(g))
            out.add(t if y or t >= d else date(yil + 1, int(a), int(g)))
        except ValueError:
            pass
    for g, ay in re.findall(r"\b(\d{1,2}) (" + "|".join(AYLAR) + r")\b", k):
        try:
            t = date(d.year, AYLAR.index(ay) + 1, int(g))
            out.add(t if t >= d else date(d.year + 1, t.month, t.day))
        except ValueError:
            pass
    if "bugun" in k:
        out.add(d)
    if "yarin" in k:
        out.add(d + timedelta(days=1))
    for n, birim in re.findall(r"\b(\d+) (is gunu|gun|hafta)\b", k):
        if birim == "is gunu":
            t, kalan = d, int(n)
            while kalan:
                t += timedelta(days=1)
                kalan -= t.weekday() < 5
            out.add(t)
        else:
            out.add(d + timedelta(days=int(n) * (7 if birim == "hafta" else 1)))
    pzt = d - timedelta(days=d.weekday())
    if re.search(r"\b(haftaya|gelecek hafta|onumuzdeki hafta)\b", k):
        out |= {pzt + timedelta(days=7 + i) for i in range(5)}
    if "bu hafta" in k:
        out |= {pzt + timedelta(days=i) for i in range(d.weekday(), 5)}
    for i, gun in enumerate(GUNLER):
        if re.search(rf"\b{gun}\w*", k):
            t = _sonraki_gun(d, i)
            out |= {t, t + timedelta(days=7)}
    if re.search(r"\bay (sonu|sonuna)", k):
        out.add(date(d.year, d.month, calendar.monthrange(d.year, d.month)[1]))
    if re.search(r"\b(gelecek|onumuzdeki) ay\b", k):
        y, a = (d.year + 1, 1) if d.month == 12 else (d.year, d.month + 1)
        out |= {date(y, a, i) for i in range(1, calendar.monthrange(y, a)[1] + 1)}
    return out


def maddeler(t: dict) -> list[tuple[str, dict]]:
    out = [(f"İhtiyaç {i}", m) for i, m in enumerate(t["ihtiyaclar"], 1)] + [(f"İtiraz {i}", m) for i, m in enumerate(t["itirazlar"], 1)]
    out += [(f"Rakip {i}", m) for i, m in enumerate(t["rakip_bilgisi"], 1)]
    return out + [(f"Aksiyon {i}", {"metin": a["is"], "satirlar": a["satirlar"]}) for i, a in enumerate(t["aksiyonlar"], 1)]


def denetle(ziyaretler: list[dict], sonuc: dict, rakipler: list[str]) -> list[str]:
    sorun = []
    for z in ziyaretler:
        v, t, l = z["kimlik"], sonuc.get(z["kimlik"]), z["satirlar"]
        if not t:
            sorun.append(f"{v}: özet yazılmadı")
            continue
        n = len(l)

        def kaynak(ad, ks):
            if not ks:
                sorun.append(f"{v}: {ad}: dayanak satır yok")
            if [k for k in ks if not 1 <= k <= n]:
                sorun.append(f"{v}: {ad}: notta olmayan satır {[k for k in ks if not 1 <= k <= n]}")
            return " ".join(l[k - 1] for k in ks if 1 <= k <= n)

        for ad, m in maddeler(t):
            src = kaynak(ad, m["satirlar"])
            fazla = sorted(sayilar(m["metin"]) - sayilar(src))
            if fazla:
                sorun.append(f"{v}: {ad}: dayandığı satırlarda olmayan sayı: {', '.join(fazla)}")
        f = t["firsat"]
        if f["tahmini_tutar"]:
            fazla = sorted(sayilar(f["tahmini_tutar"]) - sayilar(kaynak("Fırsat tutarı", f["satirlar"])))
            if fazla:
                sorun.append(f"{v}: Fırsat: tahmini tutardaki sayı notta yok: {', '.join(fazla)} (hesaplamayın; notta yoksa boş bırakın)")
        for i, a in enumerate(t["aksiyonlar"], 1):
            if not a["termin"]:
                continue
            td = tarih_coz(a["termin"])
            src = " ".join(l[k - 1] for k in a["satirlar"] if 1 <= k <= n)
            if not td:
                sorun.append(f"{v}: Aksiyon {i}: termin '{a['termin']}' GG.AA.YYYY biçiminde değil")
            elif td not in gecerli_terminler(src, z["tarih"]):
                sorun.append(f"{v}: Aksiyon {i}: termin {a['termin']} dayandığı satırlarda yazmıyor ve oradaki ifadeden hesaplanamıyor (notta termin yoksa boş bırakın)")
        kullanilan = {k for m in t["rakip_bilgisi"] for k in m["satirlar"]}
        rk = [katla(x) for x in rakipler] + ["rakip"]
        eksik = [i for i, s in enumerate(l, 1) if any(x and x in katla(s) for x in rk) and i not in kullanilan]
        if eksik:
            sorun.append(f"{v}: rakip geçen satır {eksik} rakip bilgisinde kullanılmamış")
    return sorun


def uret(ziyaretler: list[dict], duzeltme=None, onceki=None) -> dict:
    sistem = (BURASI / "prompt.md").read_text(encoding="utf-8")
    sonuc, hedef = {}, ziyaretler
    if duzeltme is not None:  # düzeltmede yalnız sorunlu ziyaretler yeniden yazılır
        sorunlu = {s.split(":")[0] for s in duzeltme}
        sonuc.update({a: b for a, b in (onceki or {}).items() if a not in sorunlu})
        hedef = [z for z in ziyaretler if z["kimlik"] in sorunlu]
    for z in hedef:
        d = [s for s in (duzeltme or []) if s.split(":")[0] == z["kimlik"]]
        satir = "\n".join(f"{i}. {s}" for i, s in enumerate(z["satirlar"], 1))
        mesaj = (f"<ziyaret musteri=\"{z['musteri']}\" tarih=\"{z['tarih']:%d.%m.%Y}\" gun=\"{GUNLER[z['tarih'].weekday()].capitalize()}\">\n{satir}\n</ziyaret>"
                 + duzeltme_blogu(d, (onceki or {}).get(z["kimlik"])))
        sonuc[z["kimlik"]] = llm.json_iste(sistem, mesaj, SEMA)
    return sonuc


def rapor_yaz(cikti: Path, ziyaretler, sonuc, kalan, tur) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "CRM"
    baslik(ws, ["Ziyaret", "Müşteri", "Tarih", "Özet", "Fırsat Aşaması", "Tahmini Tutar", "İhtiyaçlar", "İtirazlar", "Rakip Bilgisi", "Kontrol"],
           (7, 24, 11, 60, 16, 16, 45, 40, 40, 30))
    for z in ziyaretler:
        t = sonuc.get(z["kimlik"])
        sor = [s.split(": ", 1)[1] for s in kalan if s.startswith(z["kimlik"] + ":")]
        if not t:
            ws.append([z["kimlik"], z["musteri"], z["tarih"], "AI yanıt vermedi"])
            continue
        ws.append([z["kimlik"], z["musteri"], z["tarih"], t["ozet"], t["firsat"]["asama"], t["firsat"]["tahmini_tutar"],
                   *("\n".join(f"• {m['metin']}" for m in t[a]) for a in ("ihtiyaclar", "itirazlar", "rakip_bilgisi")), "; ".join(sor) or "Geçti"])
        ws.cell(ws.max_row, 3).number_format = "DD.MM.YYYY"
        boya(ws, [4, 5, 6, 7, 8, 9], AI_DOLGU)
        boya(ws, [10], KIRMIZI if sor else YESIL)
    son_not(ws, "Önemli: Özetleri CRM'e girmeden önce kontrol edin. Fiyat / indirim taahhüdü içeren aksiyonlar yönetici onayı gerektirir.", kalan, tur)
    hizala(ws)
    a = wb.create_sheet("Aksiyonlar")
    baslik(a, ["Ziyaret", "Müşteri", "İş", "Sorumlu", "Termin", "Uyarı", "Dayanak Satırlar"], (7, 24, 60, 9, 11, 30, 50))
    gorevler = []
    for z in ziyaretler:
        for x in (sonuc.get(z["kimlik"]) or {}).get("aksiyonlar", []):
            uyari = []
            if not x["termin"]:
                uyari.append("Termin belirsiz")
            if x["sorumlu"] == "Biz" and ONAY.search(x["is"]):
                uyari.append("Fiyat / indirim taahhüdü: yönetici onayı gerekir")
            td = tarih_coz(x["termin"])
            a.append([z["kimlik"], z["musteri"], x["is"], x["sorumlu"], td or x["termin"], "; ".join(uyari),
                      "\n".join(f"{k}. {z['satirlar'][k - 1]}" for k in x["satirlar"] if 1 <= k <= len(z["satirlar"]))])
            a.cell(a.max_row, 5).number_format = "DD.MM.YYYY"
            boya(a, [3], AI_DOLGU)
            if uyari:
                boya(a, [6], SARI)
            gorevler.append((z, x, td))
    hizala(a)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    yol = cikti.with_name(cikti.stem + "_gorevler.csv")
    with yol.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Müşteri", "Görev", "Sorumlu", "Termin", "Ziyaret Tarihi"])
        for z, x, td in gorevler:
            w.writerow([z["musteri"], x["is"], x["sorumlu"], f"{td:%d.%m.%Y}" if td else "", f"{z['tarih']:%d.%m.%Y}"])
    return yol


def calistir(girdi: Path, cikti: Path, rakipler: list[str] | None = None, tur: int = 2, evet: bool = False) -> dict:
    rakipler = rakipler or []
    ziyaretler = ziyaretleri_oku(girdi)
    print(f"[OK] {len(ziyaretler)} ziyaret · {sum(len(z['satirlar']) for z in ziyaretler)} not satırı")
    llm.onay_al(f"{len(ziyaretler)} ziyaretin notları (müşteri adı dahil; telefon / e-posta maskeli) gönderilecek.", evet)
    sonuc, kalan, yapilan = duzeltme_dongusu(lambda d, o: uret(ziyaretler, d, o), lambda x: denetle(ziyaretler, x, rakipler), tur)
    csv_yol = rapor_yaz(cikti, ziyaretler, sonuc, kalan, yapilan)
    return {"ziyaretler": ziyaretler, "sonuc": sonuc, "kalan": kalan, "tur": yapilan, "csv": csv_yol}


def main(argv: list[str] | None = None) -> int:
    ortam_hazirla()
    p = argparse.ArgumentParser(description="Müşteri görüşme notlarını CRM'e girilecek özet ve takip aksiyonlarına çevirir.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "ziyaret_notlari.txt",
                   help="Notlar (.txt/.docx); her ziyaret 'Ziyaret: Müşteri | GG.AA.YYYY' satırıyla başlar")
    p.add_argument("--rakipler", default="", help="Rakip firma adları (virgülle); geçtikleri satırlar rakip bilgisinde kullanılmalı")
    p.add_argument("--tur", type=int, default=2, help="Sorun kalırsa en fazla düzeltme turu (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "ziyaret_ozetleri.xlsx")
    p.add_argument("--evet", action="store_true", help="Veri gönderim onayını sormadan devam et")
    a = p.parse_args(argv)
    try:
        s = calistir(a.girdi, a.cikti, [x.strip() for x in a.rakipler.split(",") if x.strip()], a.tur, a.evet)
    except (llm.LLMHatasi, belge.BelgeHatasi) as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(s['sonuc'])}/{len(s['ziyaretler'])} ziyaret özeti · {sum(len(t['aksiyonlar']) for t in s['sonuc'].values())} aksiyon")
    for k in s["kalan"]:
        print(f"[!] Kalan: {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()} (+ {s['csv'].name})")
    print(f"[i] Kullanım: {llm.kullanim_ozeti()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
