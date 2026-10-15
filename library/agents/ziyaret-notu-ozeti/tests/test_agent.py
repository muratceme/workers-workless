"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import os
import re
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

NOTLAR = KLASOR / "ornek_veri" / "ziyaret_notlari.txt"
MESAJLAR = []


def m(metin, *s):
    return {"metin": metin, "satirlar": list(s)}


def ak(is_, sorumlu, termin, *s):
    return {"is": is_, "sorumlu": sorumlu, "termin": termin, "satirlar": list(s)}


def ozet(kimlik, dogru=True) -> dict:
    if kimlik == "V1":
        t = {"ozet": "Müşteri mevcut tedarikçinin uzun teslim sürelerinden memnun değil; fiyat farkına rağmen kısa teslimle değerlendirebilir.",
             "ihtiyaclar": [m("Aylık yaklaşık 40.000 adet karton koli; yılbaşından itibaren 55.000 adede çıkabilir", 3), m("Daha kısa teslim süresi", 2, 5)],
             "itirazlar": [m("Fiyat farkına itiraz; 1 haftalık teslim süresiyle değerlendirebilirler", 5)],
             "rakip_bilgisi": [m("X Ambalaj adet başına 7,40 TL fiyat vermiş (bizim liste fiyatımız 7,90 TL)", 4)],
             "firsat": {"asama": "Teklif", "tahmini_tutar": "", "satirlar": [3]},
             "aksiyonlar": [ak("3 farklı ölçüde numune gönderilmesi", "Biz", "23.10.2026", 6), ak("Güncel fiyat teklifinin iletilmesi", "Biz", "15.10.2026", 8),
                            ak("Teklifin yönetim kuruluna sunulması", "Müşteri", "31.10.2026", 7),
                            ak("1 haftalık teslim süresinin üretim planlamayla teyidi", "Biz", "", 5)]}
        if not dogru:
            t["ihtiyaclar"][0] = m("Aylık 45.000 adet koli", 3)
            t["rakip_bilgisi"] = []
            t["aksiyonlar"][1]["termin"] = "20.10.2026"
        return t
    return {"ozet": "Yeni ürün serisi için özel baskılı kutu ihtiyacı; tasarım dosyaları bekleniyor.",
            "ihtiyaclar": [m("Yeni ürün serisi için özel baskılı kutu; ilk sipariş 5.000 adet", 2), m("Yüksek baskı kalitesi", 4)],
            "itirazlar": [m("İlk sipariş için %10 indirim beklentisi", 5)], "rakip_bilgisi": [],
            "firsat": {"asama": "İhtiyaç belirleme", "tahmini_tutar": "5.000 adet", "satirlar": [2]},
            "aksiyonlar": [ak("Tasarım dosyalarının gönderilmesi", "Müşteri", "15.10.2026", 3), ak("%10 indirim talebi için yönetici onayı alınması", "Biz", "", 5)]}


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    MESAJLAR.append(kullanici)
    musteri = re.search(r'musteri="([^"]+)"', kullanici).group(1)
    return ozet("V1" if musteri.startswith("Örnek") else "V2", "<duzeltme>" in kullanici)


class KodTesti(unittest.TestCase):
    def test_okuma_ve_termin(self):
        z = agent.ziyaretleri_oku(NOTLAR)
        self.assertEqual([(x["kimlik"], x["musteri"], x["tarih"], len(x["satirlar"])) for x in z],
                         [("V1", "Örnek Gıda A.Ş.", date(2026, 10, 13), 8), ("V2", "Deneme Kozmetik Ltd.", date(2026, 10, 14), 5)])
        d = date(2026, 10, 13)  # salı
        self.assertIn(date(2026, 10, 23), agent.gecerli_terminler("haftaya cuma gününe kadar", d))
        self.assertIn(date(2026, 10, 31), agent.gecerli_terminler("ay sonuna kadar", d))
        self.assertEqual(agent.gecerli_terminler("2 gün içinde", d), {date(2026, 10, 15)})
        self.assertEqual(agent.gecerli_terminler("3 iş günü içinde", date(2026, 10, 16)), {date(2026, 10, 21)})
        self.assertIn(date(2026, 11, 5), agent.gecerli_terminler("5 Kasım'da toplantı", d))

    def test_denetim(self):
        z = agent.ziyaretleri_oku(NOTLAR)
        self.assertEqual(agent.denetle(z, {"V1": ozet("V1"), "V2": ozet("V2")}, []), [])
        s = agent.denetle(z, {"V1": ozet("V1", False), "V2": ozet("V2")}, [])
        self.assertEqual(len(s), 3, s)
        self.assertTrue(all(x.startswith("V1:") for x in s) and any("45000" in x for x in s) and any("rakip geçen satır [4]" in x for x in s)
                        and any("20.10.2026" in x for x in s))
        bos = dict(ozet("V2"), rakip_bilgisi=[])
        self.assertTrue(any("[4]" in x for x in agent.denetle(z[1:], {"V2": bos}, ["önceki tedarikçi"])))


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_dongusu(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "z.xlsx"
            s = agent.calistir(NOTLAR, cikti, evet=True)
            self.assertEqual((len(MESAJLAR), s["kalan"], s["tur"]), (3, [], 1))
            self.assertIn("Örnek Gıda", MESAJLAR[2])  # düzeltmede yalnız sorunlu ziyaret
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["CRM", "Aksiyonlar"])
            uyarilar = [r[5] for r in wb["Aksiyonlar"].iter_rows(min_row=2, values_only=True)]
            self.assertEqual(uyarilar.count("Termin belirsiz"), 1)
            self.assertIn("Termin belirsiz; Fiyat / indirim taahhüdü: yönetici onayı gerekir", uyarilar)
            self.assertEqual(len(s["csv"].read_text(encoding="utf-8-sig").splitlines()), 7)

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_duzeltme_kapali_ve_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            s = agent.calistir(NOTLAR, Path(tmp) / "z.xlsx", tur=0, evet=True)
            self.assertEqual(len(s["kalan"]), 3)
            self.assertEqual(agent.main(["--evet", "--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(agent.main(["--evet", "--girdi", str(NOTLAR.with_name("yok.txt"))]), 1)


if __name__ == "__main__":
    unittest.main()
