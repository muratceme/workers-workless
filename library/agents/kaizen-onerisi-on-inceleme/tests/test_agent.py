"""Agent testleri gerçek API çağırmaz: llm.json_iste sahte bir fonksiyonla değiştirilir."""
import copy
import json
import os
import re
import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path
from unittest import mock

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))
os.environ.setdefault("ANTHROPIC_API_KEY", "test-anahtari")  # sahte; testler API çağırmaz

import agent  # noqa: E402
import llm  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri" / "oneriler.csv"
CEVAP = json.loads((Path(__file__).parent / "sahte_cevaplar.json").read_text(encoding="utf-8"))
MESAJLAR = []


def sahte_json_iste(sistem, kullanici, sema, max_tokens=16000):
    """K-102'de metinde hiç rakam yokken ilk turda '5 dakika' kazanç uydurur."""
    MESAJLAR.append(kullanici)
    no = re.search(r'<oneri no="([^"]+)"', kullanici)[1]
    if no == "K-102" and "<duzeltme>" not in kullanici:
        return copy.deepcopy(CEVAP["K-102_hatali"])
    return copy.deepcopy(CEVAP[no])


def hazir():
    o = agent.onerileri_oku(ORNEK)
    for x in o:
        x["maskeli"] = agent.maskele(x)
    return {x["no"]: x for x in o}


class KodTesti(unittest.TestCase):
    def test_okuma_ve_maskeleme(self):
        o = hazir()
        self.assertEqual(list(o), ["K-101", "K-102", "K-103", "K-104", "K-105", "K-106"])
        self.assertEqual(agent.maskele({"oneren": "Ahmet Deneme", "metin": "Ahmet Deneme olarak öneriyorum; tel 0532 111 22 33"}),
                         "[ÖNEREN] olarak öneriyorum; tel [TELEFON]")
        self.assertTrue({D(25), D(4), D(2), D(45000), D(1)} <= agent.sayilar(o["K-101"]["metin"]))
        self.assertIn(D(2), agent.sayilar("iki kez çarpma tehlikesi"))
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "o.csv"
            y.write_text("Öneri No;Öneri Metni\nA;x\nA;y\n", encoding="utf-8")
            with self.assertRaises(llm.LLMHatasi):
                agent.onerileri_oku(y)

    def test_kazanc_hesabi(self):
        self.assertEqual(agent.yillik_kazanc(CEVAP["K-101"], D(300), 250, 2)[0], D(250000))      # 25 dk × 4 × 2 kişi × 250 gün × 300 / 60
        self.assertEqual(agent.yillik_kazanc(CEVAP["K-105"], D(300), 250, 2)[0], D(306000))      # 300 kg × 85 TL × 12
        self.assertEqual(agent.yillik_kazanc(CEVAP["K-103"], D(300), 250, 2)[0], D(40000))
        self.assertEqual(agent.yillik_kazanc(CEVAP["K-104"], D(300), 250, 2), (None, []))
        c = copy.deepcopy(CEVAP["K-101"])
        c["kazanc"][0].update(miktar="1", birim="saat", tekrar="", kisi="", periyot="vardiya")
        self.assertEqual(agent.yillik_kazanc(c, D(400), 200, 3)[0], D(240000))                  # 60 dk × 1 × 1 × 600 vardiya × 400 / 60
        c["kazanc"][0]["periyot"] = "hafta"
        self.assertEqual(agent.yillik_kazanc(c, D(400), 250, 3)[0], D(20000))                   # 50 hafta

    def test_denetim(self):
        o = hazir()
        nolar = set(o)
        self.assertEqual([agent.denetle(o[n], CEVAP[n], nolar) for n in nolar - {"K-102"}], [[]] * 5)
        self.assertEqual(agent.denetle(o["K-102"], CEVAP["K-102"], nolar), [])
        self.assertIn("miktar '5' öneri metninde yazmıyor", agent.denetle(o["K-102"], CEVAP["K-102_hatali"], nolar)[0])
        c = copy.deepcopy(CEVAP["K-105"])
        c["kazanc"][0].update(birim_fiyat="", tur="Malzeme")
        c["kazanc"].append({"tur": "Zaman", "miktar": "300", "birim": "kg", "tekrar": "", "kisi": "", "periyot": "ay", "birim_fiyat": "", "aciklama": ""})
        c.update(benzer=["K-105", "K-999"], uygulanabilirlik=7, uygulanabilirlik_gerekcesi="", uygulama_maliyeti="10.000", konu="")
        s = agent.denetle(o["K-105"], c, nolar)
        for parca in ("birim fiyat metinde yoksa", "dakika ya da saat", "Uygulama maliyeti '10.000'", "1-5", "gerekçesi boş", "K-105, K-999", "konu boş"):
            self.assertTrue(any(parca in x for x in s), parca)


class AgentTesti(unittest.TestCase):
    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_ornek_ve_duzeltme(self, _):
        MESAJLAR.clear()
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "k.xlsx"
            s = agent.calistir(ORNEK, cikti, evet=True)
            self.assertEqual((len(MESAJLAR), s["tur"], s["kalan"]), (7, 1, []))
            for m in MESAJLAR:
                for ad in ("Ahmet", "Elif", "Mehmet Kurgu", "Zeynep"):
                    self.assertNotIn(ad, m)
            ws = load_workbook(cikti)["Öneriler"]
            self.assertEqual([c.value for c in ws[1]][:7], ["Öneri No", "Öneren", "Bölüm", "Konu", "Yıllık Kazanç", "Uygulama Maliyeti", "Uygulanabilirlik"])
            self.assertEqual([ws.cell(2, j).value for j in (1, 2, 5, 6, 7, 11)], ["K-101", "Ahmet Deneme", 250000, 45000, 4, "K-104"])
            self.assertIn("25 dakika × 4 tekrar × 2 kişi × 250 (gün/yıl) × 300 TL/saat = 250.000 TL", ws.cell(2, 9).value)
            durum = {ws.cell(r, 1).value: ws.cell(r, 13).value or "" for r in range(2, 8)}
            self.assertIn("Kazanç hesaplanamadı", durum["K-102"])
            self.assertIn("İSG önerisi", durum["K-106"])
            self.assertNotIn("Kazanç hesaplanamadı", durum["K-106"])
            self.assertIn("Uygulama maliyeti yazılmamış", durum["K-105"])
            self.assertIn("Benzer öneri: K-101", durum["K-104"])

    @mock.patch.object(llm, "json_iste", side_effect=sahte_json_iste)
    def test_cli(self, _):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(agent.main(["--evet", "--saat-maliyeti", "420", "--cikti", str(Path(tmp) / "k.xlsx")]), 0)
            self.assertEqual(load_workbook(Path(tmp) / "k.xlsx")["Öneriler"].cell(2, 5).value, 350000)
            self.assertEqual(agent.main(["--saat-maliyeti", "0", "--evet"]), 2)
            self.assertEqual(agent.main(["--oneriler", str(Path(tmp) / "yok.csv"), "--evet"]), 1)


if __name__ == "__main__":
    unittest.main()
