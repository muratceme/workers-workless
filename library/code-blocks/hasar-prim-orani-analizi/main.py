"""
Hasar/Prim Oranı Analizi — Workers / Workless kod bloğu
Teknik (Risk Kabul / Underwriting) › Teknik Uzman (Underwriter)

Poliçe ve hasar verisinden müşteri, acente ve branş bazında hasar/prim (H/P) oranını hesaplar.
  Kazanılmış prim   = prim × (rapor tarihine kadar geçen gün ÷ poliçe gün sayısı)      (0 ile prim arasında)
  Gerçekleşen hasar = ödenen + muallak (bekleyen) tutar; hasar tarihi rapor tarihinden sonra olanlar alınmaz
  H/P %             = gerçekleşen hasar ÷ kazanılmış prim × 100
  Yenileme önerisi  = müşterinin H/P oranı --iyi (40) ve --kotu (80) eşiklerine göre üç gruptan biri; hasar adedi --sik-hasar (3) ve üzeriyse not
IBNR, hasar masrafları, rücu ve sovtaj gelirleri, reasürans payı ve komisyon hesaba girmez. Eşikler şirket politikasıdır; öneri karar değildir.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


POLICE_SUTUN = {"no": ("police no", "police numarasi"), "musteri": ("musteri", "sigortali", "sigortali adi"), "acente": ("acente", "acente adi"), "brans": ("brans", "urun"),
                "bas": ("baslangic", "baslangic tarihi"), "bit": ("bitis", "bitis tarihi"), "prim": ("prim", "net prim", "yazilan prim")}
HASAR_SUTUN = {"no": ("hasar no", "dosya no", "hasar dosya no"), "police": ("police no", "police numarasi"), "tarih": ("hasar tarihi", "tarih"),
               "odenen": ("odenen", "odenen tutar", "odenen hasar"), "muallak": ("muallak", "muallak tutar", "bekleyen", "muallak hasar")}
Z = Decimal(0)


def kazanilmis(prim: Decimal, bas: date, bit: date, rapor: date) -> Decimal:
    if rapor <= bas:
        return Z
    if rapor >= bit:
        return prim
    return yuvarla(prim * (rapor - bas).days / (bit - bas).days)


def oku(police_yol: Path, hasar_yol: Path, rapor: date) -> tuple[list[dict], list[dict]]:
    uy, policeler = [], {}
    for r in kayitlar(police_yol, POLICE_SUTUN, ("no", "prim", "bas", "bit")):
        no = metin(r["no"])
        if not no:
            continue
        kim = f"{no} (satır {r['_satir']})"
        prim, bas, bit = para(r["prim"]), tarih(r["bas"]), tarih(r["bit"])
        if no in policeler:
            uy.append(uyari("Yüksek", "Mükerrer poliçe", kim, "Poliçe numarası tekrar ediyor; ilk satır kullanıldı. Zeyilleri tek satırda toplayın."))
        elif prim is None:
            uy.append(uyari("Yüksek", "Prim okunamadı", kim, f"'{metin(r['prim'])}'. Poliçe analize alınmadı."))
        elif not bas or not bit or bit <= bas:
            uy.append(uyari("Yüksek", "Poliçe tarihleri okunamadı", kim, f"Başlangıç '{metin(r['bas'])}', bitiş '{metin(r['bit'])}'. Poliçe analize alınmadı."))
        else:
            policeler[no] = {"no": no, "musteri": metin(r["musteri"]) or "Belirtilmemiş", "acente": metin(r["acente"]) or "Belirtilmemiş", "brans": metin(r["brans"]) or "Belirtilmemiş",
                             "bas": bas, "bit": bit, "prim": prim, "kazanilmis": kazanilmis(prim, bas, bit, rapor), "odenen": Z, "muallak": Z, "adet": 0}
    for r in kayitlar(hasar_yol, HASAR_SUTUN, ("police",)):
        pno = metin(r["police"])
        if not pno:
            continue
        kim = f"{metin(r['no']) or 'Hasar'} (satır {r['_satir']})"
        p, t = policeler.get(pno), tarih(r["tarih"])
        od, mu = para(r["odenen"]), para(r["muallak"])
        if not p:
            uy.append(uyari("Yüksek", "Hasarın poliçesi bulunamadı", kim, f"Poliçe '{pno}' poliçe listesinde yok; hasar analize alınmadı."))
            continue
        if (od is None and metin(r["odenen"])) or (mu is None and metin(r["muallak"])):
            uy.append(uyari("Yüksek", "Hasar tutarı okunamadı", kim, f"Ödenen '{metin(r['odenen'])}', muallak '{metin(r['muallak'])}'. Okunamayan tutar 0 sayıldı."))
        if metin(r["tarih"]) and not t:
            uy.append(uyari("Orta", "Hasar tarihi okunamadı", kim, f"'{metin(r['tarih'])}'. Hasar rapor tarihinden önce sayıldı."))
        if t and t > rapor:
            uy.append(uyari("Bilgi", "Hasar rapor tarihinden sonra", kim, f"Hasar tarihi {t:%d.%m.%Y}; analize alınmadı."))
            continue
        if t and not p["bas"] <= t <= p["bit"]:
            uy.append(uyari("Orta", "Hasar tarihi poliçe süresi dışında", kim, f"Hasar {t:%d.%m.%Y}, poliçe {p['bas']:%d.%m.%Y} – {p['bit']:%d.%m.%Y}. Hasar yine de sayıldı; kontrol edin."))
        p["odenen"] += od or Z
        p["muallak"] += mu or Z
        p["adet"] += 1
    return list(policeler.values()), uy


def hp(hasar: Decimal, prim: Decimal) -> Decimal | None:
    return yuvarla(hasar / prim * 100) if prim > 0 else None


def grupla(policeler: list[dict], anahtar: str) -> list[dict]:
    g = {}
    for p in policeler:
        x = g.setdefault(p[anahtar], {"ad": p[anahtar], "police": 0, "prim": Z, "kazanilmis": Z, "odenen": Z, "muallak": Z, "adet": 0})
        x["police"] += 1
        for a in ("prim", "kazanilmis", "odenen", "muallak", "adet"):
            x[a] += p[a]
    for x in g.values():
        x["hasar"] = x["odenen"] + x["muallak"]
        x["hp"] = hp(x["hasar"], x["kazanilmis"])
    return sorted(g.values(), key=lambda x: (x["hp"] is None, -(x["hp"] or 0), x["ad"]))


def oneri(x: dict, iyi: Decimal, kotu: Decimal, sik: int) -> tuple[str, str]:
    """(grup, öneri metni)"""
    if x["hp"] is None:
        return "Veri yetersiz", "Kazanılmış prim yok; oran hesaplanamadı."
    ek = f" Hasar adedi {x['adet']}: hasar sıklığını ayrıca inceleyin." if x["adet"] >= sik else ""
    if x["hp"] <= iyi:
        return "Olumlu", f"H/P %{x['hp']}: mevcut şartlarla yenileme değerlendirilebilir.{ek}"
    if x["hp"] <= kotu:
        return "İzlenecek", f"H/P %{x['hp']}: yenilemede fiyat ve muafiyetler gözden geçirilmeli.{ek}"
    return "Olumsuz", f"H/P %{x['hp']}: şart ve fiyat revizyonu veya yenilememe değerlendirilmeli.{ek}"


def analiz(policeler: list[dict], rapor: date, iyi: Decimal = Decimal(40), kotu: Decimal = Decimal(80), sik: int = 3, yenileme_gun: int = 60) -> dict:
    for p in policeler:
        p["hasar"] = p["odenen"] + p["muallak"]
        p["hp"] = hp(p["hasar"], p["kazanilmis"])
    musteri = grupla(policeler, "musteri")
    for m in musteri:
        m["grup"], m["oneri"] = oneri(m, iyi, kotu, sik)
    md = {m["ad"]: m for m in musteri}
    yaklasan = sorted((p for p in policeler if 0 <= (p["bit"] - rapor).days <= yenileme_gun), key=lambda p: (p["bit"], p["no"]))
    for p in yaklasan:
        p["musteri_hp"], p["grup"], p["oneri"] = md[p["musteri"]]["hp"], md[p["musteri"]]["grup"], md[p["musteri"]]["oneri"]
    t = {"police": len(policeler)} | {a: sum((p[a] for p in policeler), Z) for a in ("prim", "kazanilmis", "odenen", "muallak", "hasar")} | {"adet": sum(p["adet"] for p in policeler)}
    t["hp"] = hp(t["hasar"], t["kazanilmis"])
    return {"toplam": t, "musteri": musteri, "acente": grupla(policeler, "acente"), "brans": grupla(policeler, "brans"), "policeler": policeler, "yaklasan": yaklasan,
            "esik": (iyi, kotu, sik, yenileme_gun)}


def rapor_yaz(cikti: Path, r: dict, rapor: date, uy: list[dict]) -> None:
    wb, t = Workbook(), r["toplam"]
    iyi, kotu, sik, yg = r["esik"]
    renk = {"Olumlu": YESIL, "İzlenecek": SARI, "Olumsuz": KIRMIZI, "Veri yetersiz": MAVI}

    def hp_renk(v):
        return None if v is None else YESIL if v <= iyi else SARI if v <= kotu else KIRMIZI

    ws = wb.active
    ws.title = "Özet"
    _baslik(ws, ["Gösterge", "Değer"], (38, 20))
    for ad, v in (("Rapor tarihi", f"{rapor:%d.%m.%Y}"), ("Poliçe sayısı", t["police"]), ("Yazılan prim", t["prim"]), ("Kazanılmış prim", t["kazanilmis"]), ("Hasar adedi", t["adet"]),
                  ("Ödenen hasar", t["odenen"]), ("Muallak hasar", t["muallak"]), ("Gerçekleşen hasar", t["hasar"]), ("Hasar / prim oranı %", t["hp"]),
                  ("Eşikler", f"olumlu ≤ %{iyi}, olumsuz > %{kotu}, sık hasar ≥ {sik} adet")):
        _satir(ws, [ad, v], para_s=(2,) if isinstance(v, Decimal) else ())
    ws.append([])
    ws.append(["H/P = (ödenen + muallak hasar) ÷ kazanılmış prim. IBNR, hasar masrafı, rücu-sovtaj, reasürans ve komisyon dahil değildir."])
    bas = ["Poliçe", "Yazılan Prim", "Kazanılmış Prim", "Hasar Adedi", "Ödenen", "Muallak", "Gerçekleşen Hasar", "H/P %"]
    ws = wb.create_sheet("Müşteri Bazında")
    _baslik(ws, ["Müşteri", *bas, "Grup", "Yenileme Önerisi"], (32, 8, 16, 16, 9, 16, 16, 16, 10, 14, 80))
    for x in r["musteri"]:
        _satir(ws, [x["ad"], x["police"], x["prim"], x["kazanilmis"], x["adet"], x["odenen"], x["muallak"], x["hasar"], x["hp"], x["grup"], x["oneri"]], para_s=(3, 4, 6, 7, 8, 9),
               dolgu={9: hp_renk(x["hp"]), 10: renk[x["grup"]]}, sar=(11,))
    for ad, veri, kol in (("Acente Bazında", r["acente"], "Acente"), ("Branş Bazında", r["brans"], "Branş")):
        ws = wb.create_sheet(ad)
        _baslik(ws, [kol, *bas], (30, 8, 16, 16, 9, 16, 16, 16, 10))
        for x in veri:
            _satir(ws, [x["ad"], x["police"], x["prim"], x["kazanilmis"], x["adet"], x["odenen"], x["muallak"], x["hasar"], x["hp"]], para_s=(3, 4, 6, 7, 8, 9), dolgu={9: hp_renk(x["hp"])})
    ws = wb.create_sheet("Yenilemesi Yaklaşanlar")
    _baslik(ws, ["Poliçe No", "Müşteri", "Acente", "Branş", "Bitiş", "Prim", "Poliçe H/P %", "Müşteri H/P %", "Grup", "Öneri", "Karar (siz doldurun)"], (14, 30, 22, 18, 12, 16, 11, 11, 14, 70, 24))
    for p in r["yaklasan"]:
        _satir(ws, [p["no"], p["musteri"], p["acente"], p["brans"], p["bit"], p["prim"], p["hp"], p["musteri_hp"], p["grup"], p["oneri"], ""], para_s=(6, 7, 8), tarih_s=(5,),
               dolgu={9: renk[p["grup"]]}, sar=(10,))
    ws = wb.create_sheet("Poliçeler")
    _baslik(ws, ["Poliçe No", "Müşteri", "Acente", "Branş", "Başlangıç", "Bitiş", "Yazılan Prim", "Kazanılmış Prim", "Hasar Adedi", "Ödenen", "Muallak", "H/P %"],
            (14, 30, 22, 18, 12, 12, 16, 16, 9, 16, 16, 10))
    for p in r["policeler"]:
        _satir(ws, [p["no"], p["musteri"], p["acente"], p["brans"], p["bas"], p["bit"], p["prim"], p["kazanilmis"], p["adet"], p["odenen"], p["muallak"], p["hp"]], para_s=(7, 8, 10, 11, 12),
               tarih_s=(5, 6), dolgu={12: hp_renk(p["hp"])})
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(police_yol: Path, hasar_yol: Path, cikti: Path, rapor: date, **esik) -> dict:
    policeler, uy = oku(police_yol, hasar_yol, rapor)
    if not policeler:
        raise ValueError(f"{police_yol.name}: analiz edilecek poliçe yok.")
    r = analiz(policeler, rapor, **esik)
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, rapor, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Müşteri, acente ve branş bazında hasar/prim oranını hesaplar; yenileme önerisi çıkarır.")
    p.add_argument("--policeler", type=Path, default=ORNEK / "policeler.csv", help="Poliçe No, Müşteri, Acente, Branş, Başlangıç, Bitiş, Prim")
    p.add_argument("--hasarlar", type=Path, default=ORNEK / "hasarlar.csv", help="Hasar No, Poliçe No, Hasar Tarihi, Ödenen, Muallak")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 01.10.2026)")
    p.add_argument("--iyi", default="40", help="Olumlu sayılacak H/P üst sınırı %% (varsayılan 40)")
    p.add_argument("--kotu", default="80", help="Olumsuz sayılacak H/P alt sınırı %% (varsayılan 80)")
    p.add_argument("--sik-hasar", type=int, default=3, help="Sık hasar notu için asgari hasar adedi (varsayılan 3)")
    p.add_argument("--yenileme-gun", type=int, default=60, help="Yenilemesi yaklaşan sayılacak gün (varsayılan 60)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "hasar_prim_analizi.xlsx")
    a = p.parse_args(argv)
    for y in (a.policeler, a.hasarlar):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    rapor = tarih(a.tarih) if a.tarih else (date(2026, 10, 1) if a.policeler == ORNEK / "policeler.csv" else date.today())
    iyi, kotu = para(a.iyi), para(a.kotu)
    if not rapor or iyi is None or kotu is None or not 0 <= iyi < kotu or a.sik_hasar < 1 or a.yenileme_gun < 0:
        print("[X] --tarih GG.AA.YYYY biçiminde; --iyi < --kotu; --sik-hasar en az 1 olmalı")
        return 2
    try:
        r = calistir(a.policeler, a.hasarlar, a.cikti, rapor, iyi=iyi, kotu=kotu, sik=a.sik_hasar, yenileme_gun=a.yenileme_gun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {t['police']} poliçe · kazanılmış prim {t['kazanilmis']:,.2f} · gerçekleşen hasar {t['hasar']:,.2f} · H/P %{t['hp'] if t['hp'] is not None else '—'}")
    for m in r["musteri"]:
        if m["grup"] == "Olumsuz":
            print(f"[!] {m['ad']}: {m['oneri']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
