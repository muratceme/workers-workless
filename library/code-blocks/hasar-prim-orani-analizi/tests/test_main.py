import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
RAPOR = date(2026, 10, 1)


def pol(no, musteri="M1", kazanilmis=1000, odenen=0, muallak=0, adet=0, acente="A", brans="Yangın", bit=date(2027, 6, 1), prim=None):
    return {"no": no, "musteri": musteri, "acente": acente, "brans": brans, "bas": date(2026, 6, 1), "bit": bit, "prim": D(prim if prim is not None else kazanilmis),
            "kazanilmis": D(kazanilmis), "odenen": D(odenen), "muallak": D(muallak), "adet": adet}


class HesapTesti(unittest.TestCase):
    def test_kazanilmis_prim(self):
        k = lambda bas, bit: main.kazanilmis(D(73000), bas, bit, RAPOR)   # noqa: E731
        self.assertEqual(k(date(2026, 4, 1), date(2027, 4, 1)), D("36600.00"))               # 183 ÷ 365
        self.assertEqual(k(date(2025, 10, 1), date(2026, 10, 1)), D(73000))                  # süresi dolmuş
        self.assertEqual(k(date(2026, 10, 1), date(2027, 10, 1)), D(0))                      # bugün başlıyor
        self.assertEqual(k(date(2026, 10, 15), date(2027, 10, 15)), D(0))
        self.assertEqual(k(date(2026, 9, 30), date(2027, 9, 30)), D("200.00"))               # 1 ÷ 365

    def test_oran_ve_gruplar(self):
        r = main.analiz([pol("P1", "M1", 1000, 300, 100, 2), pol("P2", "M1", 1000, 0, 0, 0, brans="Nakliyat"), pol("P3", "M2", 500, 450, 0, 1, acente="B"), pol("P4", "M3", 0, 0, 0, 0, prim=900)], RAPOR)
        self.assertEqual([(m["ad"], m["kazanilmis"], m["hasar"], m["hp"], m["grup"]) for m in r["musteri"]],
                         [("M2", D(500), D(450), D("90.00"), "Olumsuz"), ("M1", D(2000), D(400), D("20.00"), "Olumlu"), ("M3", D(0), D(0), None, "Veri yetersiz")])
        self.assertEqual({a["ad"]: a["hp"] for a in r["acente"]}, {"A": D("20.00"), "B": D("90.00")})
        self.assertEqual({b["ad"]: (b["police"], b["hp"]) for b in r["brans"]}, {"Yangın": (3, D("56.67")), "Nakliyat": (1, D("0.00"))})
        t = r["toplam"]
        self.assertEqual((t["police"], t["prim"], t["kazanilmis"], t["hasar"], t["adet"], t["hp"]), (4, D(3400), D(2500), D(850), 3, D("34.00")))

    def test_oneri_esikleri(self):
        o = lambda hp, adet=0, **kw: main.oneri({"hp": None if hp is None else D(hp), "adet": adet}, kw.get("iyi", D(40)), kw.get("kotu", D(80)), 3)   # noqa: E731
        self.assertEqual([o(x)[0] for x in (0, 40, "40.01", 80, "80.01", 250, None)], ["Olumlu", "Olumlu", "İzlenecek", "İzlenecek", "Olumsuz", "Olumsuz", "Veri yetersiz"])
        self.assertEqual(o(50, iyi=D(60))[0], "Olumlu")
        self.assertIn("Hasar adedi 3", o(10, 3)[1])
        self.assertNotIn("Hasar adedi", o(10, 2)[1])

    def test_yenilemesi_yaklasan(self):
        l = [pol("P1", bit=date(2026, 10, 1)), pol("P2", bit=date(2026, 11, 30)), pol("P3", bit=date(2026, 12, 1)), pol("P4", bit=date(2026, 9, 30)), pol("P5", "M2", 100, 90, bit=date(2026, 10, 20))]
        r = main.analiz(l, RAPOR)
        self.assertEqual([(p["no"], p["grup"]) for p in r["yaklasan"]], [("P1", "Olumlu"), ("P5", "Olumsuz"), ("P2", "Olumlu")])
        self.assertEqual([p["no"] for p in main.analiz(l, RAPOR, yenileme_gun=61)["yaklasan"]], ["P1", "P5", "P2", "P3"])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "h.xlsx"
            r = main.calistir(ORNEK / "policeler.csv", ORNEK / "hasarlar.csv", cikti, RAPOR)
            t = r["toplam"]
            self.assertEqual((t["police"], t["kazanilmis"], t["hasar"], t["adet"], t["hp"]), (10, D("951500.00"), D(352000), 10, D("36.99")))
            self.assertEqual({m["ad"].split()[0] + " " + m["ad"].split()[1]: (m["hp"], m["grup"]) for m in r["musteri"]},
                             {"Kurgu Mobilya": (D("122.95"), "Olumsuz"), "Deneme Plastik": (D("52.00"), "İzlenecek"), "Sahte Tekstil": (D("30.77"), "Olumlu"),
                              "Örnek Otomotiv": (D("28.13"), "Olumlu"), "Hayali Lojistik": (D("0.00"), "Olumlu"), "Uydurma Gıda": (D("0.00"), "Olumlu"), "Yeni Müşteri": (None, "Veri yetersiz")})
            self.assertEqual({a["ad"][-1]: a["hp"] for a in r["acente"]}, {"A": D("61.06"), "B": D("25.23"), "C": D("27.09")})
            self.assertEqual({b["ad"]: b["hp"] for b in r["brans"]}, {"Makine Kırılması": D("69.90"), "Nakliyat": D("65.79"), "Yangın": D("29.90")})
            self.assertEqual([p["no"] for p in r["yaklasan"]], ["P-101", "P-102", "P-106", "P-108", "P-109"])
            self.assertIn("Hasar adedi 4", next(m["oneri"] for m in r["musteri"] if m["ad"].startswith("Örnek Otomotiv")))
            self.assertEqual([u["tur"] for u in r["uyarilar"]], ["Hasarın poliçesi bulunamadı", "Hasar rapor tarihinden sonra"])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Müşteri Bazında", "Acente Bazında", "Branş Bazında", "Yenilemesi Yaklaşanlar", "Poliçeler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            p, h = Path(tmp) / "p.csv", Path(tmp) / "h.csv"
            p.write_text("Poliçe No;Müşteri;Başlangıç;Bitiş;Prim\nP1;M;01.10.2025;01.10.2026;1.000,00\nP1;M;01.10.2025;01.10.2026;5\nP2;M;01.10.2025;01.10.2026;çok\n"
                         "P3;M;01.10.2026;01.10.2025;10\n", encoding="utf-8")
            h.write_text("Hasar No;Poliçe No;Hasar Tarihi;Ödenen;Muallak\nH1;P1;01.01.2025;100;x\nH2;P1;dün;50;\nH3;P9;;1;\n", encoding="utf-8")
            pl, uy = main.oku(p, h, RAPOR)
            self.assertEqual([(x["no"], x["odenen"], x["muallak"], x["adet"], x["acente"]) for x in pl], [("P1", D(150), D(0), 2, "Belirtilmemiş")])
            self.assertEqual(sorted(u["tur"] for u in uy), ["Hasar tarihi okunamadı", "Hasar tarihi poliçe süresi dışında", "Hasar tutarı okunamadı", "Hasarın poliçesi bulunamadı",
                                                            "Mükerrer poliçe", "Poliçe tarihleri okunamadı", "Prim okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--policeler", str(p), "--hasarlar", str(h), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "01.10.2026"]), 0)
            self.assertEqual(main.main(arg + ["--tarih", "dün"]), 2)
            self.assertEqual(main.main(arg + ["--iyi", "90"]), 2)
            self.assertEqual(main.main(["--policeler", str(Path(tmp) / "yok.csv")]), 1)
            p.write_text("Poliçe No;Başlangıç;Bitiş;Prim\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
