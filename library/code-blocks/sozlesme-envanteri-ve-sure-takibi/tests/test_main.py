import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)


def s(karsi, bitis, otomatik=False, bildirim=None, ay=None, sorumlu="S", tur="T", tutar=None, no=""):
    return {"no": no, "karsi": karsi, "konu": "", "tur": tur, "bas": None, "bitis": bitis, "otomatik": otomatik, "ay": ay, "bildirim": bildirim, "sorumlu": sorumlu, "tutar": None if tutar is None else D(tutar), "kim": f"{no} · {karsi}".strip(" ·")}


def tek(*a, gun=90, **kw):
    r = main.degerlendir([s("K", *a, **kw)], BUGUN, gun)
    x = r["sozlesmeler"][0]
    return x["durum"], x["guncel"], x["bildirim_son"], x["karar"], x["kalan"], sorted(u["tur"] for u in r["uyarilar"])


class SureTesti(unittest.TestCase):
    def test_ay_ekle(self):
        self.assertEqual([main.ay_ekle(date(2025, 8, 31), n) for n in (12, 6, 1, 18)], [date(2026, 8, 31), date(2026, 2, 28), date(2025, 9, 30), date(2027, 2, 28)])
        self.assertEqual((main.ay_ekle(date(2024, 2, 29), 12), main.ay_ekle(date(2026, 2, 28), 12), main.ay_ekle(date(2026, 11, 15), 3), main.ay_ekle(date(2026, 1, 30), 1)),
                         (date(2025, 2, 28), date(2027, 2, 28), date(2027, 2, 15), date(2026, 2, 28)))

    def test_otomatik_yenilenmeyen(self):
        self.assertEqual(tek(date(2026, 10, 11)), ("Süresi dolmuş", date(2026, 10, 11), None, date(2026, 10, 11), -1, ["Süresi dolmuş sözleşme"]))
        self.assertEqual([tek(b)[::4] for b in (date(2026, 10, 12), date(2027, 1, 10), date(2027, 1, 11))], [("Bitiş yaklaşıyor", 0), ("Bitiş yaklaşıyor", 90), ("Yürürlükte", 91)])
        self.assertEqual(tek(date(2027, 1, 11), bildirim=60)[:4], ("Yürürlükte", date(2027, 1, 11), None, date(2027, 1, 11)))          # otomatik değilse bildirim süresi karar tarihini değiştirmez
        self.assertEqual((tek(date(2026, 11, 30), gun=30)[0], tek(None, bildirim=30), tek(None)[5]), ("Yürürlükte", ("Süresiz", None, None, None, None, []), ["Süresiz sözleşmede fesih bildirim süresi yok"]))

    def test_otomatik_yenilenen(self):
        self.assertEqual(tek(date(2026, 11, 30), True, 30), ("Fesih bildirimi yaklaşıyor", date(2026, 11, 30), date(2026, 10, 31), date(2026, 10, 31), 19, ["Fesih bildirimi için son gün yaklaşıyor"]))
        self.assertEqual(tek(date(2026, 12, 31), True, 90), ("Bildirim süresi geçti", date(2026, 12, 31), date(2026, 10, 2), date(2026, 10, 2), -10, ["Fesih bildirim süresi geçti"]))
        self.assertEqual([tek(date(2026, 12, 31), True, g)[::4] for g in (80, 81)], [("Fesih bildirimi yaklaşıyor", 0), ("Bildirim süresi geçti", -1)])
        self.assertEqual(tek(date(2025, 8, 31), True, 60), ("Yürürlükte", date(2027, 8, 31), date(2027, 7, 2), date(2027, 7, 2), 263, ["Otomatik yenilenmiş sayıldı"]))
        self.assertEqual(tek(date(2026, 4, 30), True, 10, ay=6)[:5], ("Fesih bildirimi yaklaşıyor", date(2026, 10, 31), date(2026, 10, 21), date(2026, 10, 21), 9))
        self.assertEqual(tek(date(2026, 10, 12), True, 0)[:5], ("Fesih bildirimi yaklaşıyor", date(2026, 10, 12), date(2026, 10, 12), date(2026, 10, 12), 0))
        self.assertEqual(tek(date(2026, 12, 31), True), ("Fesih bildirimi yaklaşıyor", date(2026, 12, 31), date(2026, 12, 31), date(2026, 12, 31), 80, ["Fesih bildirim süresi yok", "Fesih bildirimi için son gün yaklaşıyor"]))

    def test_ozet_ve_sira(self):
        l = [s("Uzak", date(2028, 1, 1), tur="Kira", sorumlu="İdari", tutar=100), s("Dolmuş", date(2026, 1, 1), tur="Hizmet", sorumlu="", tutar=50, no="A1"), s("Yakın-2", date(2026, 11, 20), tur="kira", sorumlu="idari", tutar=30),
             s("Yakın-1", date(2026, 11, 1), tur="Hizmet", sorumlu="BT", no="a1"), s("Süresiz", None, bildirim=30, tur="Tedarik", sorumlu="BT", tutar=500), s("Geçti", date(2026, 10, 20), True, 30, tur="Lisans", sorumlu="BT", tutar=70)]
        r = main.degerlendir(l, BUGUN)
        self.assertEqual([x["karsi"] for x in r["sozlesmeler"]], ["Dolmuş", "Geçti", "Yakın-1", "Yakın-2", "Uzak", "Süresiz"])
        self.assertEqual(r["say"], {"Süresi dolmuş": 1, "Bildirim süresi geçti": 1, "Fesih bildirimi yaklaşıyor": 0, "Bitiş yaklaşıyor": 2, "Yürürlükte": 1, "Süresiz": 1})
        self.assertEqual([(x["ad"], x["adet"], x["aksiyon"], x["tutar"]) for x in r["turler"]], [("Hizmet", 2, 2, D(50)), ("Kira", 2, 1, D(130)), ("Lisans", 1, 1, D(70)), ("Tedarik", 1, 0, D(500))])
        self.assertEqual([(x["ad"], x["adet"], x["aksiyon"], x["tutar"]) for x in r["sorumlular"]], [("BT", 3, 2, D(570)), ("İdari", 2, 1, D(130)), ("(sorumlu yok)", 1, 1, D(50))])
        self.assertEqual(sorted(u["tur"] for u in r["uyarilar"] if u["onem"] != "Yüksek"), ["Fesih bildirim süresi geçti", "Sorumlu yok", "Sözleşme bitişi yaklaşıyor", "Sözleşme bitişi yaklaşıyor", "Sözleşme numarası tekrar ediyor"])
        with self.assertRaises(ValueError):
            main.degerlendir([], BUGUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "s.xlsx"
            r = main.calistir(ORNEK / "sozlesmeler.csv", cikti, BUGUN)
            self.assertEqual([(x["no"], x["durum"], x["guncel"], x["karar"], x["kalan"]) for x in r["sozlesmeler"]],
                             [("S-004", "Süresi dolmuş", date(2026, 6, 30), date(2026, 6, 30), -104), ("S-001", "Bildirim süresi geçti", date(2026, 12, 31), date(2026, 10, 2), -10),
                              ("S-002", "Fesih bildirimi yaklaşıyor", date(2026, 11, 30), date(2026, 10, 31), 19), ("S-007", "Fesih bildirimi yaklaşıyor", date(2026, 12, 31), date(2026, 12, 31), 80),
                              ("S-003", "Bitiş yaklaşıyor", date(2026, 10, 31), date(2026, 10, 31), 19), ("S-008", "Yürürlükte", date(2027, 3, 1), date(2027, 3, 1), 140),
                              ("S-005", "Yürürlükte", date(2027, 8, 31), date(2027, 7, 2), 263), ("S-006", "Süresiz", None, None, None)])
            self.assertEqual([(u["onem"], u["tur"], u["kim"].split(" · ")[0]) for u in r["uyarilar"]],
                             [("Yüksek", "Fesih bildirimi için son gün yaklaşıyor", "S-002"), ("Yüksek", "Fesih bildirimi için son gün yaklaşıyor", "S-007"), ("Yüksek", "Süresi dolmuş sözleşme", "S-004"),
                              ("Orta", "Fesih bildirim süresi geçti", "S-001"), ("Orta", "Sözleşme bitişi yaklaşıyor", "S-003"), ("Bilgi", "Fesih bildirim süresi yok", "S-007"), ("Bilgi", "Otomatik yenilenmiş sayıldı", "S-005")])
            self.assertEqual([(x["ad"], x["adet"], x["aksiyon"], x["tutar"]) for x in r["sorumlular"]],
                             [("İdari İşler", 3, 2, D(3480000)), ("Bilgi Teknolojileri", 1, 1, D(650000)), ("Mali İşler", 2, 1, D(580000)), ("Lojistik", 1, 1, D(0)), ("Satın Alma", 1, 0, D(0))])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Envanter", "Tür Özeti", "Sorumlu Özeti", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "s.csv"
            y.write_text("No;Firma;Sözleşme Türü;İmza Tarihi;Bitiş Tarihi;Otomatik Yenilenir;Yenileme Dönemi;İhbar Süresi;Sözleşme Bedeli\n1;A;Kira;01.01.2025;31.12.2026;EVET;24;90;1.250.000,50\n2;B;;;;hayır;0;-5;x\n"
                         "3;C;;01.01.2026;x;;;;\n4;D;;01.01.2026;01.01.2025;;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["karsi"], x["tur"], x["bitis"], x["otomatik"], x["ay"], x["bildirim"], x["tutar"]) for x in l], [u["tur"] for u in uy]),
                             ([("1", "A", "Kira", date(2026, 12, 31), True, 24, 90, D("1250000.50")), ("2", "B", "(tür yok)", None, False, None, None, None)], ["Sözleşme okunamadı"] * 2))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--sozlesmeler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "12.10.2026", "--gun", "120"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual((main.main(arg + ["--gun", "-1"]), main.main(arg + ["--tarih", "x"])), (2, 2))
            self.assertEqual(main.main(["--sozlesmeler", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Karşı Taraf;Konu\nA;x\n", "Karşı Taraf;Bitiş\nA;x\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
