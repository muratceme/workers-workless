"""
Sözleşme Envanteri ve Süre Takibi — Gerisi Bende kod bloğu
Hukuk › Hukuk Müdürü

Sözleşmelerin bitiş, yenileme ve fesih bildirim tarihlerini izler; karar verilmesi gereken son günü öne çıkarır.
  Güncel bitiş          = otomatik yenilenen sözleşmenin bitişi geçmişse, yenileme süresi (boşsa 12 ay) eklenerek --tarih sonrasına taşınır
  Fesih bildirim son günü = güncel bitiş − fesih bildirim süresi (takvim günü)
  Karar tarihi          = otomatik yenilenenlerde fesih bildirim son günü (süre yazılmamışsa bitiş); diğerlerinde bitiş
  Durum                 = Süresi dolmuş          : otomatik yenilenmeyen ve bitişi geçmiş
                          Bildirim süresi geçti  : otomatik yenilenen ve bildirim son günü geçmiş (sözleşme yenilenecek)
                          Fesih bildirimi yaklaşıyor / Bitiş yaklaşıyor : karar tarihine --gun (90) gün veya daha az kalmış
                          Yürürlükte · Süresiz (bitiş tarihi yok)
Tarihler sizin girdiğiniz verilere dayanır. Sürenin nasıl hesaplanacağı ve bildirimin şekli sözleşme metnine bağlıdır; metni esas alın.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
import calendar

SUTUN = {"no": ("sozlesme no", "no", "sozlesme numarasi"), "karsi": ("karsi taraf", "taraf", "firma"), "konu": ("konu", "sozlesme konusu", "aciklama"), "tur": ("tur", "sozlesme turu", "kategori"),
         "bas": ("baslangic", "baslangic tarihi", "imza tarihi"), "bitis": ("bitis", "bitis tarihi", "sona erme tarihi"), "otomatik": ("otomatik yenileme", "otomatik yenilenir", "yenileme"),
         "ay": ("yenileme suresi", "yenileme suresi ay", "yenileme donemi"), "bildirim": ("fesih bildirim suresi", "fesih bildirim suresi gun", "bildirim suresi", "ihbar suresi"),
         "sorumlu": ("sorumlu", "sorumlu birim", "sozlesme sahibi"), "tutar": ("yillik tutar", "tutar", "sozlesme bedeli")}
EVET = ("evet", "e", "x", "1", "var", "true", "otomatik")
SIRA = {"Süresi dolmuş": 0, "Bildirim süresi geçti": 1, "Fesih bildirimi yaklaşıyor": 2, "Bitiş yaklaşıyor": 3, "Yürürlükte": 4, "Süresiz": 5}
D_RENK = {"Süresi dolmuş": KIRMIZI, "Bildirim süresi geçti": SARI, "Fesih bildirimi yaklaşıyor": KIRMIZI, "Bitiş yaklaşıyor": SARI, "Yürürlükte": YESIL}
Z = Decimal(0)


def ay_ekle(t: date, n: int) -> date:
    y, a = divmod(t.year * 12 + t.month - 1 + n, 12)
    son = calendar.monthrange(y, a + 1)[1]
    return date(y, a + 1, son if t.day == calendar.monthrange(t.year, t.month)[1] else min(t.day, son))


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("karsi", "bitis")):
        karsi, no = metin(r["karsi"]), metin(r["no"])
        if not karsi and not no:
            continue
        kim = f"{no} · {karsi}".strip(" ·")
        bas, bit = tarih(r["bas"]), tarih(r["bitis"])
        if (metin(r["bitis"]) and bit is None) or (bas and bit and bit < bas):
            uy.append(uyari("Yüksek", "Sözleşme okunamadı", f"{kim} (satır {r['_satir']})", f"Başlangıç '{metin(r['bas'])}', bitiş '{metin(r['bitis'])}'. Bitiş okunabilir ve başlangıçtan sonra olmalı. Sözleşme envantere alınmadı."))
            continue
        ay, bildirim, tutar = para(r["ay"]), para(r["bildirim"]), para(r["tutar"])
        l.append({"no": no, "karsi": karsi, "konu": metin(r["konu"]), "tur": metin(r["tur"]) or "(tür yok)", "bas": bas, "bitis": bit, "otomatik": katla(r["otomatik"]) in EVET, "ay": int(ay) if ay and ay >= 1 else None,
                  "bildirim": int(bildirim) if bildirim is not None and bildirim >= 0 else None, "sorumlu": metin(r["sorumlu"]), "tutar": tutar if tutar and tutar > 0 else None, "kim": kim})
    return l, uy


def _ozet(sozlesmeler: list[dict], alan: str) -> list[dict]:
    o = {}
    for s in sozlesmeler:
        x = o.setdefault(katla(s[alan]), {"ad": s[alan] or "(sorumlu yok)", "adet": 0, "aksiyon": 0, "tutar": Z})
        x["adet"] += 1
        x["aksiyon"] += SIRA[s["durum"]] < 4
        x["tutar"] += s["tutar"] or Z
    return sorted(o.values(), key=lambda x: (-x["aksiyon"], -x["tutar"], katla(x["ad"])))


def degerlendir(sozlesmeler: list[dict], bugun: date, gun: int = 90) -> dict:
    if not sozlesmeler:
        raise ValueError("Envanterde sözleşme yok.")
    uy, gorulen = [], set()
    for s in sozlesmeler:
        if s["no"] and katla(s["no"]) in gorulen:
            uy.append(uyari("Orta", "Sözleşme numarası tekrar ediyor", s["no"], "Aynı numara birden çok satırda var; ek protokolse ayrı numara verin."))
        gorulen.add(katla(s["no"]))
        s.update(guncel=s["bitis"], yenilenme=0, bildirim_son=None, karar=None, kalan=None)
        if s["bitis"] is None:
            s["durum"] = "Süresiz"
            if s["bildirim"] is None:
                uy.append(uyari("Bilgi", "Süresiz sözleşmede fesih bildirim süresi yok", s["kim"], "Bitiş tarihi olmayan sözleşmenin fesih bildirim süresi yazılmamış; sözleşme metninden ekleyin."))
        else:
            if s["otomatik"]:
                while s["guncel"] < bugun:
                    s["guncel"], s["yenilenme"] = ay_ekle(s["bitis"], (s["yenilenme"] + 1) * (s["ay"] or 12)), s["yenilenme"] + 1
                if s["yenilenme"]:
                    uy.append(uyari("Bilgi", "Otomatik yenilenmiş sayıldı", s["kim"], f"Yazılı bitiş {s['bitis']:%d.%m.%Y} geçmiş; {s['yenilenme']} kez {s['ay'] or 12} ay yenilendiği varsayılarak güncel bitiş {s['guncel']:%d.%m.%Y} alındı."
                                    " Yenilemenin gerçekleştiğini ve yeni koşulları doğrulayıp envanteri güncelleyin."))
                if s["bildirim"] is None:
                    uy.append(uyari("Bilgi", "Fesih bildirim süresi yok", s["kim"], "Otomatik yenilenen sözleşmede fesih bildirim süresi yazılmamış; karar tarihi olarak bitiş alındı. Sözleşme metninden ekleyin."))
                s["bildirim_son"] = s["guncel"] - timedelta(days=s["bildirim"] or 0)
            s["karar"] = s["bildirim_son"] or s["guncel"]
            s["kalan"] = (s["karar"] - bugun).days
            if s["kalan"] < 0:
                s["durum"] = "Bildirim süresi geçti" if s["otomatik"] else "Süresi dolmuş"
            elif s["kalan"] <= gun:
                s["durum"] = "Fesih bildirimi yaklaşıyor" if s["otomatik"] else "Bitiş yaklaşıyor"
            else:
                s["durum"] = "Yürürlükte"
        if s["durum"] == "Süresi dolmuş":
            uy.append(uyari("Yüksek", "Süresi dolmuş sözleşme", s["kim"], f"Bitiş {s['bitis']:%d.%m.%Y}, {-s['kalan']} gün geçmiş. İlişki sürüyorsa yenileme veya yeni sözleşme yapın; sürmüyorsa envanterden çıkarın."))
        elif s["durum"] == "Bildirim süresi geçti":
            uy.append(uyari("Orta", "Fesih bildirim süresi geçti", s["kim"], f"Bildirim için son gün {s['bildirim_son']:%d.%m.%Y} idi ({-s['kalan']} gün önce). Sözleşme {s['guncel']:%d.%m.%Y} tarihinde {s['ay'] or 12} ay daha yenilenecek görünüyor."))
        elif s["durum"] == "Fesih bildirimi yaklaşıyor":
            uy.append(uyari("Yüksek", "Fesih bildirimi için son gün yaklaşıyor", s["kim"], f"Yenilemeyi istemiyorsanız en geç {s['karar']:%d.%m.%Y} tarihine kadar ({s['kalan']} gün) bildirimde bulunun; aksi halde sözleşme "
                                                                                          f"{s['guncel']:%d.%m.%Y} tarihinde yenilenir."))
        elif s["durum"] == "Bitiş yaklaşıyor":
            uy.append(uyari("Orta", "Sözleşme bitişi yaklaşıyor", s["kim"], f"Bitiş {s['guncel']:%d.%m.%Y}, {s['kalan']} gün kaldı. Yenileme, yeni teklif veya sonlandırma kararını verin."))
        if not s["sorumlu"]:
            uy.append(uyari("Orta", "Sorumlu yok", s["kim"], "Sözleşmeye sorumlu birim veya kişi atanmamış."))
    turler, sorumlular = _ozet(sozlesmeler, "tur"), _ozet(sozlesmeler, "sorumlu")
    sozlesmeler.sort(key=lambda s: (SIRA[s["durum"]], s["karar"] or date.max, katla(s["karsi"])))
    return {"sozlesmeler": sozlesmeler, "say": {d: sum(s["durum"] == d for s in sozlesmeler) for d in SIRA}, "turler": turler, "sorumlular": sorumlular, "bugun": bugun, "gun": gun, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Envanter"
    _baslik(ws, ["Durum", "Sözleşme No", "Karşı Taraf", "Konu", "Tür", "Sorumlu", "Başlangıç", "Yazılı Bitiş", "Otomatik Yenileme", "Güncel Bitiş", "Fesih Bildirim Süresi (gün)", "Fesih Bildirim Son Günü", "Karar Tarihi", "Kalan Gün",
                 "Yıllık Tutar"], (24, 12, 26, 32, 14, 20, 12, 12, 10, 12, 11, 13, 12, 9, 16))
    for s in r["sozlesmeler"]:
        _satir(ws, [s["durum"], s["no"], s["karsi"], s["konu"], s["tur"], s["sorumlu"], s["bas"], s["bitis"], "Evet" if s["otomatik"] else "Hayır", s["guncel"], s["bildirim"], s["bildirim_son"], s["karar"], s["kalan"], s["tutar"]],
               para_s=(15,), tarih_s=(7, 8, 10, 12, 13), dolgu={1: D_RENK.get(s["durum"]), 10: MAVI if s["yenilenme"] else None}, sar=(4,))
    ws.append([])
    ws.append([f"Durum tarihi {r['bugun']:%d.%m.%Y}; karar tarihine {r['gun']} gün veya daha az kalanlar 'yaklaşıyor' sayılır. Mavi güncel bitiş, otomatik yenilemenin gerçekleştiği varsayımıyla hesaplanmıştır."])
    for ad, veri, baslik in (("Tür Özeti", r["turler"], "Tür"), ("Sorumlu Özeti", r["sorumlular"], "Sorumlu")):
        ws = wb.create_sheet(ad)
        _baslik(ws, [baslik, "Sözleşme", "Karar Bekleyen", "Yıllık Tutar"], (28, 10, 10, 18))
        for x in veri:
            _satir(ws, [x["ad"], x["adet"], x["aksiyon"] or None, x["tutar"] or None], para_s=(4,), dolgu={3: SARI if x["aksiyon"] else None})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, gun: int = 90) -> dict:
    sozlesmeler, uy = oku(yol)
    r = degerlendir(sozlesmeler, bugun, gun)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Sözleşmelerin bitiş, yenileme ve fesih bildirim tarihlerini izler.")
    p.add_argument("--sozlesmeler", type=Path, help="Sözleşme No, Karşı Taraf, Konu, Tür, Başlangıç, Bitiş, Otomatik Yenileme, Yenileme Süresi (ay), Fesih Bildirim Süresi (gün), Sorumlu, Yıllık Tutar")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--gun", type=int, default=90, help="Karar tarihine kaç gün kala uyarılsın (varsayılan 90)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "sozlesme_envanteri.xlsx")
    a = p.parse_args(argv)
    ornek = a.sozlesmeler is None
    yol = ORNEK / "sozlesmeler.csv" if ornek else a.sozlesmeler
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    if bugun is None or not 0 <= a.gun <= 730:
        print("[X] --tarih GG.AA.YYYY biçiminde, --gun 0–730 arasında olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, bugun, a.gun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(r['sozlesmeler'])} sözleşme · " + " · ".join(f"{d}: {n}" for d, n in r["say"].items() if n))
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
