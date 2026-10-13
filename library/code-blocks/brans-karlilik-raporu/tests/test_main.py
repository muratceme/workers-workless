import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def s(brans, yazilan, dev_kpk=0, son_kpk=0, odenen=0, dev_mual=0, son_mual=0, komisyon=0, gider=0, donem=""):
    return {"brans": brans, "donem": donem, "yazilan": D(yazilan), "dev_kpk": D(dev_kpk), "son_kpk": D(son_kpk), "odenen": D(odenen), "dev_mual": D(dev_mual), "son_mual": D(son_mual), "komisyon": D(komisyon), "gider": D(gider)}


def k(b):
    return (b["kp"], b["hasar"], b["hp"], b["kom_oran"], b["gider_oran"], b["bilesik"], b["teknik"], b["marj"])


class HesapTesti(unittest.TestCase):
    def test_gostergeler(self):
        r = main.hesapla([s("A", 1000, 400, 500, 500, 200, 330, 135, 90), s("B", 2000, 0, 800, 900, 100, 400, 240, 120)])
        self.assertEqual([(b["brans"], b["durum"], b["pay"]) + k(b) for b in r["branslar"]],
                         [("B", "Teknik zarar", D("66.67"), D(1200), D(1200), D("100.00"), D("20.00"), D("10.00"), D("130.00"), D(-360), D("-30.00")),
                          ("A", "Olağan", D("33.33"), D(900), D(630), D("70.00"), D("15.00"), D("10.00"), D("95.00"), D(45), D("5.00"))])
        self.assertEqual(k(r["toplam"]), (D(2100), D(1830), D("87.14"), D("17.86"), D("10.00"), D("115.00"), D(-315), D("-15.00")))
        self.assertEqual(([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]], r["donemler"]), ([("Yüksek", "Bileşik oran eşiğin üstünde", "B")], []))

    def test_esikler_ve_sinirlar(self):
        d = lambda **kw: {b["brans"]: b["durum"] for b in main.hesapla([s("H80", 100, odenen=80), s("H81", 100, odenen=81), s("B100", 100, odenen=70, komisyon=20, gider=10), s("B101", 100, odenen=70, komisyon=20, gider=11),   # noqa: E731
                                                            s("Sıfır", 100, 0, 100, 10), s("Eksi", 100, odenen=10, dev_mual=50)], **kw)["branslar"]}
        self.assertEqual(d(), {"H80": "Olağan", "H81": "Hasar oranı yüksek", "B100": "Olağan", "B101": "Teknik zarar", "Sıfır": "Oran hesaplanamadı", "Eksi": "Olağan"})
        self.assertEqual(d(hp_esik=D(60), bilesik_esik=D(105)), {"H80": "Hasar oranı yüksek", "H81": "Hasar oranı yüksek", "B100": "Hasar oranı yüksek", "B101": "Hasar oranı yüksek", "Sıfır": "Oran hesaplanamadı", "Eksi": "Olağan"})
        r = main.hesapla([s("Sıfır", 100, 0, 100, 10), s("Eksi", 100, odenen=10, dev_mual=50)])
        self.assertEqual(([(b["brans"], b["kp"], b["hasar"], b["hp"], b["bilesik"], b["teknik"]) for b in r["branslar"]], sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"])),
                         ([("Eksi", D(100), D(-40), D("-40.00"), D("-40.00"), D(140)), ("Sıfır", D(0), D(10), None, None, D(-10))], [("Bilgi", "Gerçekleşen hasar eksi", "Eksi"), ("Orta", "Kazanılmış prim sıfır ya da eksi", "Sıfır")]))
        with self.assertRaises(ValueError):
            main.hesapla([])

    def test_toplama_ve_donem(self):
        r = main.hesapla([s("Kasko", 100, odenen=50, donem="2026-02"), s("KASKO", 300, odenen=250, donem="2026-01"), s("Yangın", 200, odenen=20, donem="2026-01")])
        self.assertEqual([(b["brans"], b["yazilan"], b["hasar"], b["hp"]) for b in r["branslar"]], [("Kasko", D(400), D(300), D("75.00")), ("Yangın", D(200), D(20), D("10.00"))])
        self.assertEqual([(b["donem"], b["brans"], b["kp"], b["hp"]) for b in r["donemler"]], [("2026-01", "KASKO", D(300), D("83.33")), ("2026-01", "Yangın", D(200), D("10.00")), ("2026-02", "Kasko", D(100), D("50.00"))])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "brans_verisi.csv", cikti)
            self.assertEqual([(b["brans"], b["durum"], b["kp"], b["hasar"], b["hp"], b["bilesik"], b["teknik"]) for b in r["branslar"]],
                             [("Trafik", "Teknik zarar", D(56000000), D(60800000), D("108.57"), D("128.57"), D(-16000000)), ("Kasko", "Olağan", D(45000000), D(31500000), D("70.00"), D("95.00"), D(2250000)),
                              ("Sağlık", "Teknik zarar", D(32000000), D(27200000), D("85.00"), D("103.00"), D(-960000)), ("Yangın", "Olağan", D(20000000), D(8000000), D("40.00"), D("70.00"), D(6000000)),
                              ("Nakliyat", "Olağan", D(8000000), D(2800000), D("35.00"), D("63.00"), D(2960000))])
            t = r["toplam"]
            self.assertEqual((t["kp"], t["hasar"], t["komisyon"], t["gider"], t["teknik"], t["hp"], t["bilesik"]), (D(161000000), D(130300000), D(20990000), D(15460000), D(-5750000), D("80.93"), D("103.57")))
            self.assertEqual([(u["onem"], u["kim"]) for u in r["uyarilar"]], [("Yüksek", "Sağlık"), ("Yüksek", "Trafik")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Branşlar", "Hesap Ayrıntısı", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "b.csv"
            y.write_text("Ürün;Yıl;Yazılan Primler;KPK;Ödenen Tazminat;Muallak;Komisyon\nA;2026;1.000,50;200;300;;50\n;2026;5;;;;\nB;2026;x;;;;\nC;2025;100;;y;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["brans"], x["donem"], x["yazilan"], x["son_kpk"], x["odenen"], x["son_mual"], x["komisyon"], x["gider"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("A", "2026", D("1000.50"), D(200), D(300), D(0), D(50), D(0))], [("Satır okunamadı", "B"), ("Satır okunamadı", "C")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--veri", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--hp", "75", "--bilesik", "98,5"]), 0)
            self.assertEqual(load_workbook(c).sheetnames, ["Branşlar", "Hesap Ayrıntısı", "Dönemler", "Uyarılar"])
            self.assertEqual([main.main(arg + v) for v in (["--hp", "0"], ["--bilesik", "x"])], [2] * 2)
            self.assertEqual(main.main(["--veri", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Branş;Prim\nA;5\n", "Branş;Yazılan Prim\n;5\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
