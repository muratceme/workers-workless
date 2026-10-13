"""
Branş Kârlılık Raporu — Workers / Workless kod bloğu
Teknik (Risk Kabul / Underwriting) › Teknik Müdür

Branş bazında yazılan prim, karşılık değişimleri, hasar ve giderlerden teknik kârlılık göstergelerini üretir.
  Kazanılmış prim     = yazılan prim + devreden kazanılmamış primler karşılığı − dönem sonu kazanılmamış primler karşılığı
  Gerçekleşen hasar   = ödenen hasar + dönem sonu muallak − devreden muallak
  Hasar prim oranı %  = gerçekleşen hasar ÷ kazanılmış prim × 100
  Komisyon oranı %    = komisyon gideri ÷ kazanılmış prim × 100
  Gider oranı %       = faaliyet gideri ÷ kazanılmış prim × 100
  Bileşik oran %      = (gerçekleşen hasar + komisyon + faaliyet gideri) ÷ kazanılmış prim × 100
  Teknik sonuç        = kazanılmış prim − gerçekleşen hasar − komisyon − faaliyet gideri;   marj % = teknik sonuç ÷ kazanılmış prim
Tutarları brüt ya da reasürans payı düşülmüş (net) verebilirsiniz; hepsi aynı esasta olmalıdır. Yatırım geliri hesaba girmez.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"brans": ("brans", "urun", "brans adi"), "donem": ("donem", "yil", "ay"), "yazilan": ("yazilan prim", "brut yazilan prim", "yazilan primler"),
         "dev_kpk": ("devreden kpk", "devreden kazanilmamis primler karsiligi", "donem basi kpk"), "son_kpk": ("donem sonu kpk", "kazanilmamis primler karsiligi", "kpk"),
         "odenen": ("odenen hasar", "odenen hasarlar", "odenen tazminat"), "dev_mual": ("devreden muallak", "donem basi muallak", "devreden muallak hasar karsiligi"),
         "son_mual": ("donem sonu muallak", "muallak hasar karsiligi", "muallak"), "komisyon": ("komisyon gideri", "komisyon", "uretim komisyonu"), "gider": ("faaliyet gideri", "genel gider", "faaliyet giderleri")}
TUTAR = ("yazilan", "dev_kpk", "son_kpk", "odenen", "dev_mual", "son_mual", "komisyon", "gider")


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("brans", "yazilan")):
        ad = metin(r["brans"])
        if not ad:
            continue
        d = {a: para(r[a]) if metin(r[a]) else Decimal(0) for a in TUTAR}
        if any(v is None for v in d.values()):
            uy.append(uyari("Yüksek", "Satır okunamadı", ad, f"Satır {r['_satir']}: " + ", ".join(f"{SUTUN[a][0]} '{metin(r[a])}'" for a in TUTAR if d[a] is None) + " okunamadı. Satır alınmadı."))
            continue
        l.append({"brans": ad, "donem": metin(r["donem"]), **d})
    return l, uy


def _oran(pay: Decimal, payda: Decimal) -> Decimal | None:
    return yuvarla(pay * 100 / payda) if payda > 0 else None


def gosterge(t: dict) -> dict:
    """t: TUTAR alanlarının toplamı. Kazanılmış prim sıfır ya da eksiyse oranlar hesaplanmaz."""
    kp = t["yazilan"] + t["dev_kpk"] - t["son_kpk"]
    hasar = t["odenen"] + t["son_mual"] - t["dev_mual"]
    teknik = kp - hasar - t["komisyon"] - t["gider"]
    return {**t, "kp": kp, "hasar": hasar, "hp": _oran(hasar, kp), "kom_oran": _oran(t["komisyon"], kp), "gider_oran": _oran(t["gider"], kp), "bilesik": _oran(hasar + t["komisyon"] + t["gider"], kp),
            "teknik": teknik, "marj": _oran(teknik, kp)}


def _topla(satirlar: list[dict]) -> dict:
    return {a: sum((s[a] for s in satirlar), Decimal(0)) for a in TUTAR}


def hesapla(satirlar: list[dict], hp_esik: Decimal = Decimal(80), bilesik_esik: Decimal = Decimal(100)) -> dict:
    if not satirlar:
        raise ValueError("Raporlanacak branş satırı yok.")
    grup, donem_grup = {}, {}
    for s in satirlar:
        grup.setdefault(katla(s["brans"]), []).append(s)
        donem_grup.setdefault((s["donem"], katla(s["brans"])), []).append(s)
    toplam = gosterge(_topla(satirlar))
    branslar = sorted(({"brans": g[0]["brans"], **gosterge(_topla(g))} for g in grup.values()), key=lambda b: (-b["yazilan"], katla(b["brans"])))
    uy = []
    for b in branslar:
        b["pay"] = _oran(b["yazilan"], toplam["yazilan"])
        b["durum"] = "Oran hesaplanamadı" if b["bilesik"] is None else "Teknik zarar" if b["bilesik"] > bilesik_esik else "Hasar oranı yüksek" if b["hp"] > hp_esik else "Olağan"
        if b["bilesik"] is None:
            uy.append(uyari("Orta", "Kazanılmış prim sıfır ya da eksi", b["brans"], f"Kazanılmış prim {b['kp']:,.2f}; oranlar hesaplanmadı. Karşılık tutarlarını kontrol edin."))
        elif b["durum"] == "Teknik zarar":
            uy.append(uyari("Yüksek", "Bileşik oran eşiğin üstünde", b["brans"], f"Bileşik oran %{b['bilesik']} (eşik %{bilesik_esik}): hasar %{b['hp']}, komisyon %{b['kom_oran']}, gider %{b['gider_oran']}. Teknik sonuç {b['teknik']:,.2f}."))
        elif b["durum"] == "Hasar oranı yüksek":
            uy.append(uyari("Orta", "Hasar prim oranı eşiğin üstünde", b["brans"], f"Hasar prim oranı %{b['hp']} (eşik %{hp_esik}); bileşik oran %{b['bilesik']}."))
        if b["hasar"] < 0:
            uy.append(uyari("Bilgi", "Gerçekleşen hasar eksi", b["brans"], f"Muallak karşılıktaki azalış ödenen hasardan büyük ({b['hasar']:,.2f}). Önceki dönem karşılığı fazla ayrılmış olabilir."))
    donemler = []
    if any(s["donem"] for s in satirlar):
        donemler = [{"donem": d, "brans": g[0]["brans"], **gosterge(_topla(g))} for (d, _), g in sorted(donem_grup.items(), key=lambda x: (x[0][0], x[0][1]))]
    return {"branslar": branslar, "toplam": toplam, "donemler": donemler, "uyarilar": uy}


BASLIK = ["Yazılan Prim", "Kazanılmış Prim", "Gerçekleşen Hasar", "Komisyon", "Faaliyet Gideri", "Teknik Sonuç", "Hasar Prim %", "Komisyon %", "Gider %", "Bileşik %", "Teknik Marj %"]


def _deger(b: dict) -> list:
    return [b["yazilan"], b["kp"], b["hasar"], b["komisyon"], b["gider"], b["teknik"], b["hp"], b["kom_oran"], b["gider_oran"], b["bilesik"], b["marj"]]


def rapor_yaz(cikti: Path, r: dict, bilesik_esik: Decimal) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Branşlar"
    _baslik(ws, ["Branş", "Durum", "Prim Payı %"] + BASLIK, (24, 20, 9) + (17,) * 6 + (11,) * 5, "B2")
    for b in r["branslar"]:
        _satir(ws, [b["brans"], b["durum"], b["pay"]] + _deger(b), para_s=range(3, 15),
               dolgu={2: {"Teknik zarar": KIRMIZI, "Hasar oranı yüksek": SARI, "Oran hesaplanamadı": SARI, "Olağan": YESIL}[b["durum"]], 9: KIRMIZI if b["teknik"] < 0 else None})
    n = _satir(ws, ["TOPLAM", None, None] + _deger(r["toplam"]), para_s=range(3, 15))
    for h in ws[n]:
        h.font = Font(bold=True)
    ws = wb.create_sheet("Hesap Ayrıntısı")
    _baslik(ws, ["Branş", "Yazılan Prim", "Devreden KPK", "Dönem Sonu KPK", "Kazanılmış Prim", "Ödenen Hasar", "Devreden Muallak", "Dönem Sonu Muallak", "Gerçekleşen Hasar"], (24,) + (18,) * 8, "B2")
    for b in r["branslar"] + [{"brans": "TOPLAM", **r["toplam"]}]:
        _satir(ws, [b["brans"], b["yazilan"], b["dev_kpk"], b["son_kpk"], b["kp"], b["odenen"], b["dev_mual"], b["son_mual"], b["hasar"]], para_s=range(2, 10))
    if r["donemler"]:
        ws = wb.create_sheet("Dönemler")
        _baslik(ws, ["Dönem", "Branş"] + BASLIK, (12, 24) + (17,) * 6 + (11,) * 5, "C2")
        for b in r["donemler"]:
            _satir(ws, [b["donem"], b["brans"]] + _deger(b), para_s=range(3, 14), dolgu={8: KIRMIZI if b["teknik"] < 0 else None, 12: KIRMIZI if b["bilesik"] is not None and b["bilesik"] > bilesik_esik else None})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, hp_esik: Decimal = Decimal(80), bilesik_esik: Decimal = Decimal(100)) -> dict:
    satirlar, uy = oku(yol)
    r = hesapla(satirlar, hp_esik, bilesik_esik)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, bilesik_esik)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Branş bazında kazanılmış prim, gerçekleşen hasar, hasar prim oranı, bileşik oran ve teknik sonucu raporlar.")
    p.add_argument("--veri", type=Path, help="Branş, Dönem, Yazılan Prim, Devreden KPK, Dönem Sonu KPK, Ödenen Hasar, Devreden Muallak, Dönem Sonu Muallak, Komisyon Gideri, Faaliyet Gideri")
    p.add_argument("--hp", default="80", help="Hasar prim oranı uyarı eşiği, yüzde (varsayılan 80)")
    p.add_argument("--bilesik", default="100", help="Bileşik oran uyarı eşiği, yüzde (varsayılan 100)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "brans_karlilik_raporu.xlsx")
    a = p.parse_args(argv)
    yol = a.veri or ORNEK / "brans_verisi.csv"
    hp, bil = para(a.hp), para(a.bilesik)
    if hp is None or hp <= 0 or bil is None or bil <= 0:
        print("[X] --hp ve --bilesik sıfırdan büyük yüzde olmalı (ör. 80, 100)")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, hp, bil)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {len(r['branslar'])} branş · kazanılmış prim {t['kp']:,.2f} · gerçekleşen hasar {t['hasar']:,.2f} · teknik sonuç {t['teknik']:,.2f}"
          + (f" · hasar prim %{t['hp']} · bileşik %{t['bilesik']}" if t["bilesik"] is not None else ""))
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
