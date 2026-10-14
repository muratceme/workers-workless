import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 14)


class SureTesti(unittest.TestCase):
    def test_son_tarih(self):
        # 2024/07: 10.11.2024 pazar → 11.11.2024 (gelir); kurumlar 14.11.2024 (tebliğ duyurusundaki örnek)
        self.assertEqual((main.son_tarih(2024, 7, "Gelir"), main.son_tarih(2024, 7, "Kurumlar")), (date(2024, 11, 11), date(2024, 11, 14)))
        # Aralık: gelir vergisi nisan 10, kurumlar mayıs 14
        self.assertEqual((main.son_tarih(2025, 12, "Gelir"), main.son_tarih(2025, 12, "Kurumlar")), (date(2026, 4, 10), date(2026, 5, 14)))
        # 2026/01 gelir: 10.05.2026 pazar → 11.05.2026; 2026/06 kurumlar: 14.10.2026; 2026/07 kurumlar: 14.11 cumartesi → 16.11
        self.assertEqual([main.son_tarih(2026, 1, "Gelir"), main.son_tarih(2026, 6, "Kurumlar"), main.son_tarih(2026, 7, "Kurumlar")], [date(2026, 5, 11), date(2026, 10, 14), date(2026, 11, 16)])
        # 2026/01 kurumlar: 14.05.2026 perşembe; ek tatille cuma
        self.assertEqual((main.son_tarih(2026, 1, "Kurumlar"), main.son_tarih(2026, 1, "Kurumlar", {date(2026, 5, 14)})), (date(2026, 5, 14), date(2026, 5, 15)))
        # uzatma zinciri: 11.05 → 10.06 → 30.06
        self.assertEqual(main.son_tarih(2026, 1, "Gelir", uzatma={date(2026, 5, 11): date(2026, 6, 10), date(2026, 6, 10): date(2026, 6, 30)}), date(2026, 6, 30))
        self.assertEqual([main.is_gunu(g, set()) for g in (date(2026, 10, 29), date(2026, 5, 27), date(2026, 7, 15))], [date(2026, 10, 30), date(2026, 6, 1), date(2026, 7, 16)])

    def test_donusum(self):
        self.assertEqual([main.donem_coz(x) for x in ("2026/09", "2026-9", "09.2026", "Eylül 2026", "eylul 2026", "2026/13", "x")],
                         [(2026, 9), (2026, 9), (2026, 9), (2026, 9), (2026, 9), None, None])
        self.assertEqual([main._tur(x) for x in ("GV", "Şirket", "x")], ["Gelir", "Kurumlar", None])
        self.assertEqual([main._tercih(x) for x in ("", "Geçici vergi dönemleri", "x")], ["Aylık", "Geçici vergi dönemi", None])


class TakipTesti(unittest.TestCase):
    def test_takip(self):
        m = [{"ad": "A", "vkn": "1", "tur": "Kurumlar", "tercih": "Aylık", "sorumlu": "", "anahtar": "1"}, {"ad": "B", "vkn": "2", "tur": "Gelir", "tercih": "Geçici vergi dönemi", "sorumlu": "", "anahtar": "2"}]
        y = {("1", (2026, 5)): {"olusturma": None, "imza": None, "yukleme": date(2026, 9, 1)}, ("1", (2026, 4)): {"olusturma": None, "imza": date(2026, 9, 20), "yukleme": date(2026, 9, 16)}}
        r = main.takip(m, y, BUGUN, 2026)
        d = {s["donem"]: s["durum"] for s in r["satirlar"]}
        self.assertEqual(len(d), 13)
        self.assertEqual((d["2025/12"], d["2026/04"], d["2026/05"], d["2026/06"], d["2026/09"]), ("Süresi geçti", "Geç yüklendi", "Zamanında yüklendi", "Yaklaşıyor", "Bekliyor"))
        self.assertEqual(r["ozet"], {"Süresi geçti": 7, "Yaklaşıyor": 1, "Bekliyor": 3, "Geç yüklendi": 1, "Zamanında yüklendi": 1})
        self.assertEqual(sorted({u["tur"] for u in r["uyarilar"]}), ["Berat yüklenmedi, süre geçti", "Geç yüklenmiş", "Son tarih hesaplanmadı", "Son tarih yaklaşıyor", "Tarih sırası tutarsız"])
        with self.assertRaises(ValueError):
            main.takip([], {}, BUGUN, 2026)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "mukellefler.csv", ORNEK / "yuklemeler.csv", cikti, BUGUN, 2026)
            self.assertEqual((len(r["satirlar"]), r["ozet"]["Süresi geçti"], r["ozet"]["Yaklaşıyor"], r["ozet"]["Zamanında yüklendi"]), (117, 3, 5, 81))
            self.assertEqual(r["takvim"][1][0], date(2026, 10, 12))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Dönem Takibi", "Son Tarih Takvimi", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            my, yy, ty = Path(tmp) / "m.csv", Path(tmp) / "y.csv", Path(tmp) / "t.csv"
            my.write_text("Unvan;VKN;Mükellefiyet;Tercih\nA;1;KV;\nA;1;KV;\nB;2;x;\n", encoding="utf-8")
            yy.write_text("VKN;Dönem;Yükleme Tarihi\n1;2026/01;10.05.2026\n1;2026/xx;01.01.2026\n1;2026/02;x\n", encoding="utf-8")
            ty.write_text("Tatil\n14.05.2026\n", encoding="utf-8")
            l, uy = main.mukellef_oku(my)
            self.assertEqual((len(l), len(uy)), (1, 2))
            d, uy = main.yukleme_oku(yy)
            self.assertEqual((list(d), len(uy)), ([("1", (2026, 1))], 2))
            self.assertEqual(main.tatil_oku(ty), {date(2026, 5, 14)})
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--mukellefler", str(my), "--yuklemeler", str(yy), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "14.10.2026", "--yil", "2026", "--uyari", "5", "--tatiller", str(ty), "--uzatma", "11.05.2026=30.06.2026"]), 0)
            self.assertEqual(main.main(["--mukellefler", str(my), "--cikti", c]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--uzatma", "x"], ["--uzatma", "30.06.2026=11.05.2026"], ["--yil", "1990"], ["--uyari", "-1"])], [2] * 5)
            self.assertEqual(main.main(["--mukellefler", str(Path(tmp) / "yok.csv")]), 1)
            my.write_text("Unvan;Mükellefiyet\nA;x\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
