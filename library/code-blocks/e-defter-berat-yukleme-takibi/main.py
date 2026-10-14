"""
e-Defter Berat Yükleme Takibi — Workers / Workless kod bloğu
Muhasebe ve Defter İşleri › Muhasebe Elemanı (Büro)
Mükellef bazında aylık e-defter dönemlerinin oluşturma, imzalama ve berat yükleme son tarihlerini hesaplayıp takip eder.
  Son tarih (aylık tercih) = dönemi izleyen 4. ayın 10'u (gelir vergisi mükellefi) / 14'ü (kurumlar vergisi ve diğer)
  Hesap döneminin son ayı (aralık) = yıllık beyannamenin verileceği ayı izleyen ayın 10'u (GV: nisan) / 14'ü (KV: mayıs)
  Tatil kaydırma           = son gün hafta sonu ya da resmî tatile rastlarsa izleyen ilk iş günü
  Süre uzatımı             = --uzatma "GG.AA.YYYY=GG.AA.YYYY" (GİB duyurusundaki eski → yeni son gün)
  Kaynak: Elektronik Defter Genel Tebliği (Sıra No: 1) 4.3.4 — 21.05.2024 ve 08.11.2024 değişiklikleri (güncel metinden doğrulayın)
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
MUKELLEF = {"ad": ("mukellef", "unvan", "firma", "mukellef adi"), "vkn": ("vkn", "vkn tckn", "vergi no", "tckn"), "tur": ("vergi turu", "mukellefiyet", "tur"),
            "tercih": ("yukleme tercihi", "tercih", "berat tercihi"), "sorumlu": ("sorumlu", "personel")}
YUKLEME = {"vkn": ("vkn", "vkn tckn", "vergi no", "tckn"), "ad": ("mukellef", "unvan", "firma"), "donem": ("donem", "ay"), "olusturma": ("olusturma tarihi", "olusturma"),
           "imza": ("imza tarihi", "imzalama tarihi", "imza"), "yukleme": ("berat yukleme tarihi", "yukleme tarihi", "berat tarihi", "yukleme")}
SABIT_TATIL = ((1, 1), (4, 23), (5, 1), (5, 19), (7, 15), (8, 30), (10, 29))
DINI_TATIL = {date(2026, 3, 20), date(2026, 3, 21), date(2026, 3, 22), date(2026, 5, 27), date(2026, 5, 28), date(2026, 5, 29), date(2026, 5, 30)}
AYLAR = ("Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık")


def tatil_mi(g: date, ek: set[date]) -> bool:
    return g.weekday() >= 5 or (g.month, g.day) in SABIT_TATIL or g in DINI_TATIL or g in ek


def is_gunu(g: date, ek: set[date]) -> date:
    while tatil_mi(g, ek):
        g += timedelta(days=1)
    return g


def _ay(yil: int, ay: int, ileri: int) -> tuple[int, int]:
    y, m = divmod(ay - 1 + ileri, 12)
    return yil + y, m + 1


def son_tarih(yil: int, ay: int, tur: str, ek: set[date] = frozenset(), uzatma: dict | None = None) -> date:
    gun = 10 if tur == "Gelir" else 14
    ileri = 5 if (ay == 12 and tur == "Kurumlar") else 4
    y, m = _ay(yil, ay, ileri)
    g = date(y, m, gun)
    g = g if uzatma and g in uzatma else is_gunu(g, ek)
    seen = set()
    while uzatma and g in uzatma and g not in seen:
        seen.add(g)
        g = uzatma[g]
    return g


def _tur(x) -> str | None:
    k = katla(x)
    if k in ("gelir", "gelir vergisi", "gv", "sahis", "gercek kisi"):
        return "Gelir"
    if k in ("kurumlar", "kurumlar vergisi", "kv", "sirket", "tuzel kisi", "diger"):
        return "Kurumlar"
    return None


def _tercih(x) -> str | None:
    k = katla(x)
    if not k or k in ("aylik", "ay"):
        return "Aylık"
    if "gecici" in k or k in ("3 aylik", "uc aylik"):
        return "Geçici vergi dönemi"
    return None


def donem_coz(x) -> tuple[int, int] | None:
    if isinstance(x, (date, datetime)):
        return x.year, x.month
    s = metin(x)
    m = re.fullmatch(r"(\d{4})\s*[-/.]\s*(\d{1,2})", s) or None
    if m:
        y, a = int(m.group(1)), int(m.group(2))
    else:
        m = re.fullmatch(r"(\d{1,2})\s*[-/.]\s*(\d{4})", s)
        if m:
            a, y = int(m.group(1)), int(m.group(2))
        else:
            k = katla(s).split()
            if len(k) == 2 and k[1].isdigit():
                ad = [katla(a) for a in AYLAR]
                if k[0] not in ad:
                    return None
                a, y = ad.index(k[0]) + 1, int(k[1])
            else:
                return None
    return (y, a) if 1 <= a <= 12 and 2000 <= y <= 2100 else None


def mukellef_oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, MUKELLEF, ("ad", "tur")):
        ad = metin(r["ad"])
        if not ad:
            continue
        tur, tercih = _tur(r["tur"]), _tercih(r["tercih"])
        if tur is None or tercih is None:
            uy.append(uyari("Yüksek", "Mükellef satırı okunamadı", ad, f"Satır {r['_satir']}: vergi türü ('{metin(r['tur'])}' — Gelir / Kurumlar) ya da yükleme tercihi okunamadı."))
            continue
        anahtar = metin(r["vkn"]) or katla(ad)
        if anahtar in gorulen:
            uy.append(uyari("Orta", "Mükellef iki kez yazılmış", ad, f"Satır {r['_satir']}: ikinci satır alınmadı."))
            continue
        gorulen.add(anahtar)
        l.append({"ad": ad, "vkn": metin(r["vkn"]), "tur": tur, "tercih": tercih, "sorumlu": metin(r["sorumlu"]), "anahtar": anahtar})
    return l, uy


def yukleme_oku(yol: Path) -> tuple[dict, list[dict]]:
    uy, d = [], {}
    for r in kayitlar(yol, YUKLEME, ("donem", "yukleme")):
        kim = metin(r["vkn"]) or metin(r["ad"])
        if not kim and not metin(r["donem"]):
            continue
        dn = donem_coz(r["donem"])
        t = {a: (tarih(r[a]) if metin(r[a]) else None) for a in ("olusturma", "imza", "yukleme")}
        if not kim or dn is None or any(metin(r[a]) and v is None for a, v in t.items()):
            uy.append(uyari("Yüksek", "Yükleme satırı okunamadı", kim or f"Satır {r['_satir']}", f"Satır {r['_satir']}: mükellef, dönem ('{metin(r['donem'])}') ya da tarihlerden biri okunamadı."))
            continue
        for anahtar in {metin(r["vkn"]), katla(r["ad"])} - {""}:
            d[(anahtar, dn)] = t
    return d, uy


def takip(mukellefler: list[dict], yuklemeler: dict, bugun: date, yil: int, uyari_gun: int = 10, ek: set[date] = frozenset(), uzatma: dict | None = None) -> dict:
    if not mukellefler:
        raise ValueError("Okunabilen mükellef yok.")
    uy, satirlar = [], []
    for m in mukellefler:
        if m["tercih"] != "Aylık":
            uy.append(uyari("Orta", "Son tarih hesaplanmadı", m["ad"], "Geçici vergi dönemi tercihli mükellefin son tarihleri bu sürümde hesaplanmaz; tebliğdeki süreyi elle takip edin."))
            continue
        donemler = [(yil - 1, a) for a in range(1, 13) if son_tarih(yil - 1, a, m["tur"], ek, uzatma) >= date(yil, 1, 1)]
        donemler += [(yil, a) for a in range(1, 13) if date(yil, a, 1) <= bugun.replace(day=1) - timedelta(days=1)]
        for dy, ay in donemler:
            st = son_tarih(dy, ay, m["tur"], ek, uzatma)
            t = yuklemeler.get((m["vkn"], (dy, ay))) or yuklemeler.get((katla(m["ad"]), (dy, ay))) or {"olusturma": None, "imza": None, "yukleme": None}
            kalan = (st - bugun).days
            if t["yukleme"]:
                durum = "Zamanında yüklendi" if t["yukleme"] <= st else "Geç yüklendi"
            elif kalan < 0:
                durum = "Süresi geçti"
            elif kalan <= uyari_gun:
                durum = "Yaklaşıyor"
            else:
                durum = "Bekliyor"
            s = {"ad": m["ad"], "vkn": m["vkn"], "tur": m["tur"], "sorumlu": m["sorumlu"], "donem": f"{dy}/{ay:02d}", "son": st, "kalan": kalan, **t, "durum": durum}
            satirlar.append(s)
            kim = f"{m['ad']} · {s['donem']}"
            if durum == "Süresi geçti":
                uy.append(uyari("Yüksek", "Berat yüklenmedi, süre geçti", kim, f"Son tarih {st:%d.%m.%Y}; {-kalan} gün geçti. Yaptırım riskini ve yapılacak işlemi mali müşavirle değerlendirin."))
            elif durum == "Yaklaşıyor":
                uy.append(uyari("Orta", "Son tarih yaklaşıyor", kim, f"Son tarih {st:%d.%m.%Y} ({kalan} gün)." + ("" if t["imza"] else " Defter henüz imzalanmamış.")))
            elif durum == "Geç yüklendi":
                uy.append(uyari("Bilgi", "Geç yüklenmiş", kim, f"Yükleme {t['yukleme']:%d.%m.%Y}, son tarih {st:%d.%m.%Y}."))
            if t["yukleme"] and t["imza"] and t["imza"] > t["yukleme"]:
                uy.append(uyari("Bilgi", "Tarih sırası tutarsız", kim, "İmza tarihi berat yükleme tarihinden sonra görünüyor."))
    satirlar.sort(key=lambda s: ({"Süresi geçti": 0, "Yaklaşıyor": 1, "Bekliyor": 2, "Geç yüklendi": 3, "Zamanında yüklendi": 4}[s["durum"]], s["son"], katla(s["ad"])))
    ozet = {d: sum(1 for s in satirlar if s["durum"] == d) for d in ("Süresi geçti", "Yaklaşıyor", "Bekliyor", "Geç yüklendi", "Zamanında yüklendi")}
    takvim = {}
    for s in satirlar:
        if s["durum"] in ("Yaklaşıyor", "Bekliyor", "Süresi geçti"):
            takvim.setdefault(s["son"], []).append(s)
    return {"satirlar": satirlar, "ozet": ozet, "takvim": sorted(takvim.items()), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Dönem Takibi"
    _baslik(ws, ["Mükellef", "VKN / TCKN", "Vergi Türü", "Sorumlu", "Dönem", "Son Tarih", "Kalan Gün", "Oluşturma", "İmza", "Berat Yükleme", "Durum"], (28, 13, 10, 14, 9, 11, 8, 11, 11, 12, 18), "B2")
    renk = {"Süresi geçti": KIRMIZI, "Yaklaşıyor": SARI, "Geç yüklendi": MAVI, "Zamanında yüklendi": YESIL}
    for s in r["satirlar"]:
        _satir(ws, [s["ad"], s["vkn"], s["tur"], s["sorumlu"], s["donem"], s["son"], s["kalan"], s["olusturma"], s["imza"], s["yukleme"], s["durum"]], tarih_s=(6, 8, 9, 10),
               dolgu={11: renk.get(s["durum"])})
    ws = wb.create_sheet("Son Tarih Takvimi")
    _baslik(ws, ["Son Tarih", "Mükellef Sayısı", "Mükellef / Dönem"], (11, 10, 90))
    for g, l in r["takvim"]:
        _satir(ws, [g, len(l), ", ".join(f"{s['ad']} ({s['donem']})" for s in l)], tarih_s=(1,), sar=(3,))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def tatil_oku(yol: Path) -> set[date]:
    s = set()
    for r in tablo_oku(yol):
        for c in r:
            g = tarih(c)
            if g:
                s.add(g)
    return s


def calistir(muk_yol: Path, yuk_yol: Path | None, cikti: Path, bugun: date, yil: int, uyari_gun: int = 10, ek: set[date] = frozenset(), uzatma: dict | None = None) -> dict:
    mk, uy1 = mukellef_oku(muk_yol)
    yk, uy2 = yukleme_oku(yuk_yol) if yuk_yol else ({}, [])
    r = takip(mk, yk, bugun, yil, uyari_gun, ek, uzatma)
    r["uyarilar"] = sirala(uy1 + uy2 + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="e-Defter dönemlerinin berat yükleme son tarihlerini hesaplayıp mükellef bazında takip eder.")
    p.add_argument("--mukellefler", type=Path, help="Mükellef, VKN, Vergi Türü (Gelir / Kurumlar), Yükleme Tercihi (Aylık / Geçici vergi dönemi), Sorumlu")
    p.add_argument("--yuklemeler", type=Path, help="VKN ya da Mükellef, Dönem (2026/09), Oluşturma, İmza, Berat Yükleme Tarihi")
    p.add_argument("--tarih", help="Rapor günü GG.AA.YYYY (varsayılan bugün; örnek veride 14.10.2026)")
    p.add_argument("--yil", type=int, help="Takip edilen yıl (varsayılan rapor gününün yılı)")
    p.add_argument("--uyari", type=int, default=10, help="Son tarihe kaç gün kala uyarılır (varsayılan 10)")
    p.add_argument("--tatiller", type=Path, help="Ek tatil günleri listesi (.csv/.xlsx; her hücrede bir tarih) — dini bayramlar 2026 dışında buraya")
    p.add_argument("--uzatma", action="append", default=[], help="Süre uzatımı: ESKİ=YENİ (GG.AA.YYYY=GG.AA.YYYY); birden fazla verilebilir")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "e_defter_berat_takibi.xlsx")
    a = p.parse_args(argv)
    ornek = a.mukellefler is None
    muk = a.mukellefler or ORNEK / "mukellefler.csv"
    yuk = a.yuklemeler or (ORNEK / "yuklemeler.csv" if ornek else None)
    bugun = tarih(a.tarih) if a.tarih else (date(2026, 10, 14) if ornek else date.today())
    uzatma = {}
    for u in a.uzatma:
        e, _, y = u.partition("=")
        ge, gy = tarih(e.strip()), tarih(y.strip())
        if ge is None or gy is None or gy <= ge:
            print(f"[X] --uzatma '{u}' ESKİ=YENİ biçiminde olmalı (yeni tarih eskisinden sonra)")
            return 2
        uzatma[ge] = gy
    if bugun is None or a.uyari < 0 or (a.yil is not None and not 2000 <= a.yil <= 2100):
        print("[X] --tarih GG.AA.YYYY, --uyari 0 ya da üstü, --yil 2000-2100 olmalı")
        return 2
    for y in (muk, yuk, a.tatiller):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        ek = tatil_oku(a.tatiller) if a.tatiller else set()
        r = calistir(muk, yuk, a.cikti, bugun, a.yil or bugun.year, a.uyari, ek, uzatma)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {len(r['satirlar'])} mükellef-dönem · süresi geçen {o['Süresi geçti']} · yaklaşan {o['Yaklaşıyor']} · bekleyen {o['Bekliyor']} · zamanında {o['Zamanında yüklendi']} · geç {o['Geç yüklendi']}")
    for g, l in r["takvim"][:5]:
        print(f"     {g:%d.%m.%Y}: {len(l)} mükellef-dönem")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
