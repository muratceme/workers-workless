import csv
import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "h.xlsx"
        cls.s = main.calistir(ORNEK / "satinalma.csv", cls.cikti)
        cls.u = {(u["tur"], u["kim"]) for u in cls.s["uyarilar"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_toplam_bagimsiz_hesap(self):
        toplam = D(0)
        with open(ORNEK / "satinalma.csv", encoding="utf-8") as f:
            for r in csv.DictReader(f, delimiter=";"):
                tutar = D(r["Tutar"].replace(",", "."))
                if r["Döviz"] == "TRY":
                    toplam += tutar
                elif r["Kur"]:
                    toplam += tutar * D(r["Kur"].replace(",", "."))
        self.assertEqual(self.s["toplam"], toplam)
        self.assertIn(("Kur yok", "EUR"), self.u)

    def test_abc_ve_birlestirme(self):
        k = {x["ad"]: x["sinif"] for x in self.s["kat_abc"]}
        self.assertEqual((k["Hammadde - Çelik"], k["Ambalaj"], k["Lojistik hizmeti"], k["Ofis malzemeleri"]), ("A", "A", "B", "C"))
        self.assertIn(("Tedarikçi adı birleştirildi", "Örnek Çelik A.Ş."), self.u)
        self.assertNotIn("ÖRNEK ÇELİK ANONİM ŞİRKETİ", {x["ad"] for x in self.s["ted_abc"]})
        self.assertIn(("Kategorisiz harcama", "Sınıflandırılmamış"), self.u)

    def test_oneriler(self):
        o = {(x["tur"], x["kapsam"]): x["oncelik"] for x in self.s["oneriler"]}
        self.assertEqual(o[("Fiyat birliği / toplu pazarlık", "Hammadde - Çelik")], 1)
        self.assertEqual(o[("Tek kaynak bağımlılığı", "Temizlik hizmeti")], 1)
        self.assertEqual(o[("Tedarikçi konsolidasyonu", "Yedek parça")], 3)
        self.assertEqual(o[("Fiyat artışı incelemesi", "PP granül (Hammadde - Plastik)")], 2)
        artanlar = {y["malzeme"] for y in self.s["fa"] if y["esik_ustu"]}
        self.assertEqual(artanlar, {"PP granül"})                 # genel enflasyon üzerindeki tek artış

    def test_excel(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Özet", "Pazarlık Öncelikleri", "Kategori ABC", "Tedarikçi ABC", "Kategori × Tedarikçi", "Fiyat Farkları",
                                         "Fiyat Artışları", "Aylık Trend", "Veri", "Uyarılar"])


class KuralTesti(unittest.TestCase):
    def test_abc_sinirlari(self):
        # 50 / 30 / 15 / 5: önceki kümülatif 0 → A, 50 → A, 80 → B (80'in altında değil), 95 → C
        r = main.abc({"a": D(50), "b": D(30), "c": D(15), "d": D(5)}, D("0.80"), D("0.95"))
        self.assertEqual([x["sinif"] for x in r], ["A", "A", "B", "C"])

    def test_fiyat_farki(self):
        def alim(ted, fiyat, miktar, gun=5):
            return main.Alim(1, date(2026, 3, gun), ted, main.tedarikci_anahtari(ted), "K", "Vida", D(miktar), "adet", D(fiyat), D(fiyat) * D(miktar), "TRY", "")
        ff = main.fiyat_farklari([alim("X", 10, 100), alim("Y", 12, 50), alim("X", 10, 50, 20)])
        x = ff[0]
        self.assertEqual((x["min"], x["max"], x["alim"]), (D(10), D(12), 3))
        self.assertEqual(x["ort"], D("10.5"))                             # (1000 + 600 + 500) / 200
        self.assertEqual(x["tasarruf_ort"], D(75))                         # (12 − 10,5) × 50
        self.assertEqual(x["tasarruf_min"], D(100))                        # (12 − 10) × 50
        nisan = alim("Y", 12, 50)
        nisan.tarih = date(2026, 4, 1)
        self.assertEqual(main.fiyat_farklari([alim("X", 10, 100), nisan]), [])                 # farklı ay karşılaştırılmaz
        self.assertEqual(len(main.fiyat_farklari([alim("X", 10, 100), nisan], "ceyrek")), 0)   # farklı çeyrek de

    def test_tedarikci_anahtari(self):
        self.assertEqual(main.tedarikci_anahtari("Örnek Çelik A.Ş."), main.tedarikci_anahtari("ÖRNEK ÇELİK ANONİM ŞİRKETİ"))
        self.assertEqual(main.tedarikci_anahtari("Anadolu Metal San. ve Tic. Ltd. Şti."), "anadolu metal")

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx"), "--fiyat-donem", "ceyrek"]), 0)
            self.assertEqual(main.main(["--a-esik", "96"]), 2)
            self.assertEqual(main.main(["--veri", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
