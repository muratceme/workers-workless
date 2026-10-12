"""
Kampanya Etki Analizi — Workers / Workless kod bloğu
Satın Alma ve Kategori Yönetimi › Kategori Uzmanı

Kampanya dönemindeki ek satışı (uplift), kanibalizasyonu ve kampanya kârlılığını hesaplar.
  Beklenen adet   = baz dönem adedi × kampanya gün sayısı ÷ baz dönem gün sayısı   (kampanya olmasaydı satılacak miktar)
  Ek satış        = kampanya dönemi adedi − beklenen adet;   ek satış % = ek satış ÷ beklenen × 100
  Kâr farkı       = kampanya adedi × (uygulanan fiyat − birim maliyet) − beklenen adet × (normal fiyat − birim maliyet)
  İndirim maliyeti= beklenen adet × (normal fiyat − kampanya fiyatı)               (zaten satılacak ürüne verilen indirim)
  Başabaş adet    = beklenen kâr ÷ (kampanya fiyatı − birim maliyet), yukarı yuvarlanır
  Kanibalizasyon %= aynı kategorideki kampanya dışı ürünlerin adet kaybı ÷ kampanya ürünlerinin ek satışı × 100
  Net etki        = kampanya ürünlerinin kâr farkı + aynı kategorideki diğer ürünlerin kâr farkı − kampanya gideri
Beklenen adet düz ortalamadır; mevsimsellik, stok yokluğu ve kampanya sonrası düşüş hesaba katılmaz.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"urun": ("urun", "urun adi", "sku", "stok kodu"), "kategori": ("kategori",), "rol": ("kampanyada", "kampanyada mi", "kampanya urunu", "kampanya"),
         "baz": ("baz donem adet", "baz donem adedi", "baz adet", "onceki donem adet"), "kamp": ("kampanya donemi adet", "kampanya donemi adedi", "kampanya adet", "kampanya adedi"),
         "normal": ("normal fiyat", "liste fiyati", "fiyat"), "kfiyat": ("kampanya fiyati", "indirimli fiyat"), "maliyet": ("birim maliyet", "maliyet")}
EVET = {"e", "evet", "x", "1", "var", "kampanyada", "true"}
Z = Decimal(0)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("urun", "baz", "kamp", "normal", "maliyet")):
        u = metin(r["urun"])
        if not u:
            continue
        kim = f"{u} (satır {r['_satir']})"
        v = {a: para(r[a]) for a in ("baz", "kamp", "normal", "kfiyat", "maliyet")}
        if any(v[a] is None or v[a] < 0 for a in ("baz", "kamp", "normal", "maliyet")):
            uy.append(uyari("Yüksek", "Değer okunamadı", kim, "Baz dönem adedi, kampanya dönemi adedi, normal fiyat veya birim maliyet okunamadı. Ürün analize alınmadı."))
            continue
        kampanyada = katla(r["rol"]) in EVET
        if kampanyada:
            if v["kfiyat"] is None or v["kfiyat"] < 0:
                uy.append(uyari("Orta", "Kampanya fiyatı yok", kim, f"Kampanya ürünü ama kampanya fiyatı '{metin(r['kfiyat'])}'. Normal fiyat kullanıldı."))
                v["kfiyat"] = v["normal"]
            elif v["kfiyat"] > v["normal"]:
                uy.append(uyari("Orta", "Kampanya fiyatı normal fiyattan yüksek", kim, f"Kampanya fiyatı {v['kfiyat']:f} > normal fiyat {v['normal']:f}."))
            if v["kfiyat"] <= v["maliyet"]:
                uy.append(uyari("Yüksek", "Kampanya fiyatı maliyetin altında", kim, f"Kampanya fiyatı {v['kfiyat']:f}, birim maliyet {v['maliyet']:f}: satılan her adet zarar yazar. Fiyatların KDV esasını kontrol edin."))
        if v["baz"] == 0:
            uy.append(uyari("Bilgi", "Baz dönemde satış yok", kim, "Beklenen adet 0 alındı; ek satış yüzdesi hesaplanmadı. Yeni ürün veya stoksuz dönem olabilir."))
        l.append({"urun": u, "kategori": metin(r["kategori"]) or "Belirtilmemiş", "kampanyada": kampanyada} | v)
    return l, uy


def analiz(urunler: list[dict], baz_gun: Decimal, kamp_gun: Decimal, gider: Decimal = Z) -> dict:
    if not urunler:
        raise ValueError("Analiz edilecek ürün yok.")
    kampli = {katla(u["kategori"]) for u in urunler if u["kampanyada"]}
    if not kampli:
        raise ValueError("Kampanya ürünü yok: 'Kampanyada' sütununda en az bir ürün Evet olmalı.")
    for u in urunler:
        kapsam = katla(u["kategori"]) in kampli
        fiyat = u["kfiyat"] if u["kampanyada"] else u["normal"]
        u["tur"] = "Kampanya" if u["kampanyada"] else "Diğer (aynı kategori)" if kapsam else "Kapsam dışı"
        u["beklenen"] = yuvarla(u["baz"] * kamp_gun / baz_gun)
        u["ek"] = u["kamp"] - u["beklenen"]
        u["ek_yuzde"] = yuvarla(u["ek"] / u["beklenen"] * 100) if u["beklenen"] else None
        u["ciro"], u["beklenen_ciro"] = yuvarla(u["kamp"] * fiyat), yuvarla(u["beklenen"] * u["normal"])
        u["kar"], u["beklenen_kar"] = yuvarla(u["kamp"] * (fiyat - u["maliyet"])), yuvarla(u["beklenen"] * (u["normal"] - u["maliyet"]))
        u["kar_farki"] = u["kar"] - u["beklenen_kar"]
        u["indirim"] = u["basabas"] = None
        if u["kampanyada"]:
            u["indirim"] = yuvarla(u["beklenen"] * (u["normal"] - fiyat))
            birim = fiyat - u["maliyet"]
            u["basabas"] = None if birim <= 0 else max(0, math.ceil(u["beklenen_kar"] / birim))
    kat = {}
    for u in urunler:
        if u["tur"] == "Kapsam dışı":
            continue
        k = kat.setdefault(katla(u["kategori"]), {"kategori": u["kategori"], "kamp_ek": Z, "diger_fark": Z, "kamp_kar": Z, "diger_kar": Z})
        a = "kamp" if u["kampanyada"] else "diger"
        k[a + ("_ek" if u["kampanyada"] else "_fark")] += u["ek"]
        k[a + "_kar"] += u["kar_farki"]
    for k in kat.values():
        k["kanibal"] = yuvarla(-k["diger_fark"] / k["kamp_ek"] * 100) if k["kamp_ek"] > 0 and k["diger_fark"] < 0 else None
        k["etki"] = "Kanibalizasyon" if k["diger_fark"] < 0 else "Ek satış (halo)" if k["diger_fark"] > 0 else "—"
        k["net"] = k["kamp_kar"] + k["diger_kar"]
    kapsam = [u for u in urunler if u["tur"] != "Kapsam dışı"]
    kk, dk = sum((k["kamp_kar"] for k in kat.values()), Z), sum((k["diger_kar"] for k in kat.values()), Z)
    net = kk + dk - gider
    genel = {"kamp_ek": sum((k["kamp_ek"] for k in kat.values()), Z), "diger_fark": sum((k["diger_fark"] for k in kat.values()), Z), "kamp_kar": kk, "diger_kar": dk, "gider": gider, "net": net,
             "ek_ciro": sum((u["ciro"] - u["beklenen_ciro"] for u in kapsam), Z), "indirim": sum((u["indirim"] for u in kapsam if u["indirim"] is not None), Z),
             "sonuc": "Kârlı" if net > 0 else "Zararlı" if net < 0 else "Başabaş"}
    sira = ("Kampanya", "Diğer (aynı kategori)", "Kapsam dışı")
    return {"urunler": sorted(urunler, key=lambda u: (sira.index(u["tur"]), u["kategori"], -u["kamp"])), "kategoriler": sorted(kat.values(), key=lambda k: k["net"]), "genel": genel,
            "gun": (baz_gun, kamp_gun)}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    g, (bg, kg) = r["genel"], r["gun"]
    ws = wb.active
    ws.title = "Özet"
    _baslik(ws, ["Gösterge", "Değer", "Açıklama"], (44, 18, 90))
    for ad, d, ac in (("Kampanya ürünlerinin ek satışı (adet)", g["kamp_ek"], "Kampanya dönemi adedi − beklenen adet"),
                      ("Aynı kategorideki diğer ürünlerin adet farkı", g["diger_fark"], "Negatif: kanibalizasyon; pozitif: ek satış (halo)"),
                      ("Ek ciro", g["ek_ciro"], "Kampanya kapsamındaki kategorilerde gerçekleşen ciro − beklenen ciro"),
                      ("İndirim maliyeti", g["indirim"], "Kampanya olmasa da satılacak adede verilen indirim"),
                      ("Kampanya ürünlerinin kâr farkı", g["kamp_kar"], "Gerçekleşen brüt kâr − beklenen brüt kâr"),
                      ("Diğer ürünlerin kâr farkı", g["diger_kar"], "Aynı kategorideki kampanya dışı ürünler"),
                      ("Kampanya gideri", g["gider"], "--gider ile verdiğiniz tutar (insert, reklam, teşhir, personel)"),
                      ("NET ETKİ", g["net"], g["sonuc"])):
        n = _satir(ws, [ad, d, ac], para_s=(2,), sar=(3,))
    for c in ws[n]:
        c.font = Font(bold=True)
    ws.cell(n, 2).fill = PatternFill("solid", fgColor=YESIL if g["net"] > 0 else KIRMIZI)
    ws.append([])
    ws.append([f"Baz dönem {bg.normalize():f} gün, kampanya dönemi {kg.normalize():f} gün. Beklenen adet düz ortalamadır; mevsimsellik ve kampanya sonrası satış düşüşü hesaba katılmaz."])
    ws = wb.create_sheet("Ürün Bazında")
    _baslik(ws, ["Tür", "Kategori", "Ürün", "Baz Dönem Adet", "Beklenen Adet", "Kampanya Dönemi Adet", "Ek Satış", "Ek Satış %", "Normal Fiyat", "Uygulanan Fiyat", "Birim Maliyet", "Gerçekleşen Ciro",
                 "Beklenen Kâr", "Gerçekleşen Kâr", "Kâr Farkı", "İndirim Maliyeti", "Başabaş Adet"], (20, 18, 32, 10, 10, 11, 10, 9, 11, 11, 11, 15, 14, 14, 14, 14, 10))
    for u in r["urunler"]:
        _satir(ws, [u["tur"], u["kategori"], u["urun"], u["baz"], u["beklenen"], u["kamp"], u["ek"], u["ek_yuzde"], u["normal"], u["kfiyat"] if u["kampanyada"] else u["normal"], u["maliyet"], u["ciro"],
                    u["beklenen_kar"], u["kar"], u["kar_farki"], u["indirim"], u["basabas"] if u["kampanyada"] and u["basabas"] is not None else ("yok" if u["kampanyada"] else None)],
               para_s=(5, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16), dolgu={1: MAVI if u["kampanyada"] else None, 15: (KIRMIZI if u["kar_farki"] < 0 else YESIL) if u["tur"] != "Kapsam dışı" else None})
    ws = wb.create_sheet("Kategori Bazında")
    _baslik(ws, ["Kategori", "Kampanya Ürünleri Ek Satış", "Diğer Ürünler Adet Farkı", "Etki", "Kanibalizasyon %", "Kampanya Ürünleri Kâr Farkı", "Diğer Ürünler Kâr Farkı", "Kategori Net Kâr Farkı"],
            (22, 14, 14, 18, 14, 16, 16, 16))
    for k in r["kategoriler"]:
        _satir(ws, [k["kategori"], k["kamp_ek"], k["diger_fark"], k["etki"], k["kanibal"], k["kamp_kar"], k["diger_kar"], k["net"]], para_s=(2, 3, 5, 6, 7, 8),
               dolgu={4: SARI if k["diger_fark"] < 0 else None, 8: KIRMIZI if k["net"] < 0 else YESIL})
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, baz_gun: Decimal, kamp_gun: Decimal, gider: Decimal = Z) -> dict:
    urunler, uy = oku(yol)
    r = analiz(urunler, baz_gun, kamp_gun, gider)
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Kampanyanın ek satışını, kanibalizasyonunu ve kârlılığını hesaplar.")
    p.add_argument("--satislar", type=Path, default=ORNEK / "kampanya_satislari.csv", help="Ürün, Kategori, Kampanyada (Evet/Hayır), Baz Dönem Adet, Kampanya Dönemi Adet, Normal Fiyat, Kampanya Fiyatı, Birim Maliyet")
    p.add_argument("--baz-gun", help="Baz dönemin gün sayısı (örnek veride 28)")
    p.add_argument("--kampanya-gun", help="Kampanya döneminin gün sayısı (örnek veride 7)")
    p.add_argument("--gider", help="Kampanya gideri: insert, reklam, teşhir (varsayılan 0; örnek veride 1.000)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kampanya_etki.xlsx")
    a = p.parse_args(argv)
    ornek = a.satislar == ORNEK / "kampanya_satislari.csv"
    bg, kg, gd = para(a.baz_gun or ("28" if ornek else None)), para(a.kampanya_gun or ("7" if ornek else None)), para(a.gider or ("1000" if ornek else "0"))
    if bg is None or kg is None or bg <= 0 or kg <= 0 or gd is None or gd < 0:
        print("[X] --baz-gun ve --kampanya-gun pozitif sayı olarak verilmeli; --gider sıfır veya pozitif olmalı")
        return 2
    if not a.satislar.exists():
        print(f"[X] Dosya bulunamadı: {a.satislar}")
        return 1
    try:
        r = calistir(a.satislar, a.cikti, bg, kg, gd)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    g = r["genel"]
    print(f"[OK] Kampanya ürünleri ek satış {g['kamp_ek']:,.2f} adet · diğer ürünler {g['diger_fark']:+,.2f} adet · ek ciro {g['ek_ciro']:,.2f}")
    print(f"[{'OK' if g['net'] > 0 else '!'}] Net etki {g['net']:,.2f} ({g['sonuc']}): kampanya ürünleri {g['kamp_kar']:+,.2f}, diğer ürünler {g['diger_kar']:+,.2f}, gider −{g['gider']:,.2f}")
    for k in r["kategoriler"]:
        if k["kanibal"] is not None:
            print(f"[!] {k['kategori']}: kanibalizasyon %{k['kanibal']} (diğer ürünlerde {-k['diger_fark']:,.2f} adet kayıp)")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
