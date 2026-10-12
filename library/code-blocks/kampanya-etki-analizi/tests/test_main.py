import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def u(urun, baz, kamp, normal, maliyet, kfiyat=None, kategori="Kahve"):
    return {"urun": urun, "kategori": kategori, "kampanyada": kfiyat is not None, "baz": D(baz), "kamp": D(kamp), "normal": D(normal), "kfiyat": None if kfiyat is None else D(kfiyat), "maliyet": D(maliyet)}


class AnalizTesti(unittest.TestCase):
    def test_kampanya_urunu(self):
        x = main.analiz([u("F", 400, 260, 200, 120, 160)], D(28), D(7))["urunler"][0]
        self.assertEqual((x["beklenen"], x["ek"], x["ek_yuzde"], x["ciro"], x["beklenen_kar"], x["kar"], x["kar_farki"], x["indirim"], x["basabas"]),
                         (D("100.00"), D("160.00"), D("160.00"), D("41600.00"), D("8000.00"), D("10400.00"), D("2400.00"), D("4000.00"), 200))
        x = main.analiz([u("F", 100, 10, 200, 120, 100)], D(10), D(10))["urunler"][0]       # maliyetin altında fiyat: başabaş yok
        self.assertEqual((x["kar"], x["basabas"]), (D("-200.00"), None))
        x = main.analiz([u("F", 0, 10, 200, 120, 160)], D(28), D(7))["urunler"][0]          # baz dönemde satış yok
        self.assertEqual((x["beklenen"], x["ek_yuzde"], x["basabas"]), (D("0.00"), None, 0))
        self.assertEqual(main.analiz([u("F", 100, 30, 200, 120, 160)], D(30), D(7))["urunler"][0]["beklenen"], D("23.33"))

    def test_kanibalizasyon_halo_ve_net(self):
        l = [u("F", 400, 260, 200, 120, 160), u("T", 800, 170, 60, 45), u("G", 200, 40, 180, 130), u("S", 600, 210, 300, 230, 240, "Çay"), u("B", 120, 36, 90, 60, kategori="Çay"),
             u("X", 400, 110, 25, 18, kategori="Atıştırmalık")]
        r = main.analiz(l, D(28), D(7), D(1000))
        k = {x["kategori"]: (x["kamp_ek"], x["diger_fark"], x["etki"], x["kanibal"], x["kamp_kar"], x["diger_kar"], x["net"]) for x in r["kategoriler"]}
        self.assertEqual(k, {"Kahve": (D(160), D(-40), "Kanibalizasyon", D("25.00"), D(2400), D(-950), D(1450)), "Çay": (D(60), D(6), "Ek satış (halo)", None, D(-8400), D(180), D(-8220))})
        g = r["genel"]
        self.assertEqual((g["kamp_ek"], g["diger_fark"], g["kamp_kar"], g["diger_kar"], g["net"], g["sonuc"], g["indirim"]), (D(220), D(-34), D(-6000), D(-770), D(-7770), "Zararlı", D(13000)))
        self.assertEqual(g["ek_ciro"], D(41600 + 10200 + 7200 + 50400 + 3240) - D(20000 + 12000 + 9000 + 45000 + 2700))
        self.assertEqual({x["urun"]: x["tur"] for x in r["urunler"]}["X"], "Kapsam dışı")
        self.assertEqual([x["urun"] for x in r["urunler"]][:2], ["F", "S"])
        self.assertEqual(main.analiz([u("F", 400, 260, 200, 120, 160)], D(28), D(7), D(2400))["genel"]["sonuc"], "Başabaş")
        for kotu in ([], [u("T", 1, 1, 1, 1)]):
            with self.assertRaises(ValueError):
                main.analiz(kotu, D(28), D(7))


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "k.xlsx"
            r = main.calistir(ORNEK / "kampanya_satislari.csv", cikti, D(28), D(7), D(1000))
            self.assertEqual((r["genel"]["net"], r["genel"]["sonuc"], [x["tur"] for x in r["uyarilar"]]), (D(-7770), "Zararlı", []))
            self.assertEqual({x["urun"].split()[0]: x["basabas"] for x in r["urunler"] if x["kampanyada"]}, {"Filtre": 200, "Siyah": 1050})
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Ürün Bazında", "Kategori Bazında", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "s.csv"
            y.write_text("Ürün;Kampanyada mı?;Baz Dönem Adet;Kampanya Dönemi Adet;Normal Fiyat;Kampanya Fiyatı;Birim Maliyet\nA;Evet;100;50;10,00;;6,00\nB;E;0;5;10,00;12,00;4,00\nC;x;10;5;10,00;3,00;4,00\n"
                         "D;Hayır;çok;5;10;;4\nE;;1.000;300;10;;4\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual([(x["urun"], x["kampanyada"], x["baz"], x["kfiyat"]) for x in l], [("A", True, D(100), D("10.00")), ("B", True, D(0), D("12.00")), ("C", True, D(10), D("3.00")), ("E", False, D(1000), None)])
            self.assertEqual(sorted(x["tur"] for x in uy), ["Baz dönemde satış yok", "Değer okunamadı", "Kampanya fiyatı maliyetin altında", "Kampanya fiyatı normal fiyattan yüksek", "Kampanya fiyatı yok"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--satislar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--baz-gun", "28", "--kampanya-gun", "7", "--gider", "1.500,00"]), 0)
            self.assertEqual(main.main(arg), 2)
            self.assertEqual(main.main(arg + ["--baz-gun", "0", "--kampanya-gun", "7"]), 2)
            self.assertEqual(main.main(["--satislar", str(Path(tmp) / "yok.csv"), "--baz-gun", "28", "--kampanya-gun", "7"]), 1)
            y.write_text("Ürün;Baz Dönem Adet;Kampanya Dönemi Adet;Normal Fiyat;Birim Maliyet\nA;1;1;1;1\n", encoding="utf-8")
            self.assertEqual(main.main(arg + ["--baz-gun", "28", "--kampanya-gun", "7"]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
