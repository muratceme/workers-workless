import json
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
CAL, ISY, PAR = ORNEK / "calisanlar.csv", ORNEK / "isyeri.txt", ORNEK / "tesvik_parametreleri.json"


def kisi(**kw):
    c = {"sicil": "X", "ad": "Test", "dogum": date(1995, 1, 1), "giris": date(2026, 10, 1), "kadin": True, "cinsiyet_bilinir": True, "son6": False,
         "iskur": False, "engelli": False, "odenek": False, "son_isyeri": None, "ortalama": 10.0, "sayi": 11.0, "belge": "yok"}
    return c | kw


def isyeri(**kw):
    i = {"tarih": "15.10.2026", "tarih_v": date(2026, 10, 15), "sektor": "Diğer", "imalat": False, "ozel": "Evet", "surede": "Evet", "borc": "Yok", "kayit_disi": "Hayır"}
    return i | kw


class TesvikTesti(unittest.TestCase):
    def setUp(self):
        self.p = main.parametre_oku(PAR)

    def g10(self, **kw):
        return main.t_gecici10(kisi(**kw), isyeri(), self.p["4447_gecici_10"])

    def test_yardimcilar(self):
        self.assertEqual((main.yas(date(1997, 10, 16), date(2026, 10, 15)), main.yas(date(1997, 10, 15), date(2026, 10, 15))), (28, 29))
        self.assertEqual((main.ay_ekle(date(2026, 1, 31), 1), main.ay_ekle(date(2027, 11, 30), 3)), (date(2026, 2, 28), date(2028, 2, 29)))
        self.assertEqual([main.belge_turu(x) for x in ("", "Mesleki Yeterlilik Belgesi", "MYB", "İŞKUR Kursu", "Meslek lisesi mezunu")],
                         ["yok", "myb", "myb", "mesleki_egitim", "mesleki_egitim"])

    def test_gecici10_sureleri(self):
        self.assertEqual(self.g10()["aciklama"], "24 ay · 01.10.2026-30.09.2028")
        self.assertTrue(self.g10(iskur=True, belge="myb")["aciklama"].startswith("54 ay"))
        self.assertTrue(self.g10(kadin=False, dogum=date(1998, 1, 1), belge="mesleki_egitim")["aciklama"].startswith("36 ay"))  # 28 yaş erkek
        self.assertTrue(self.g10(kadin=False, dogum=date(1990, 1, 1), belge="myb", iskur=True)["aciklama"].startswith("30 ay"))
        self.assertTrue(self.g10(kadin=False, dogum=date(1990, 1, 1), iskur=True)["aciklama"].startswith("6 ay"))
        self.assertEqual(self.g10(kadin=False, dogum=date(1990, 1, 1))["durum"], main.DEGIL)
        self.assertEqual(self.g10(kadin=False, dogum=date(1990, 1, 1), iskur=None)["durum"], main.EKSIK)

    def test_gecici10_engeller(self):
        self.assertIn("aralığında değil", self.g10(giris=date(2027, 1, 4))["aciklama"])
        self.assertIn("18 yaşını", self.g10(dogum=date(2009, 1, 1))["aciklama"])
        self.assertIn("önceki 6 ayda", self.g10(son6=True)["aciklama"])
        self.assertIn("ilave değil", self.g10(sayi=10.0)["aciklama"])
        self.assertEqual(self.g10(ortalama=None)["durum"], main.EKSIK)
        self.assertEqual(main.t_gecici10(kisi(giris=date(2020, 1, 6)), isyeri(), self.p["4447_gecici_10"])["aciklama"], "süre 05.01.2022 tarihinde doldu")
        r = main.t_gecici10(kisi(), isyeri(borc="Var", ozel="Hayır"), self.p["4447_gecici_10"])
        self.assertEqual(r["durum"], main.DEGIL)
        self.assertIn("özel sektör", r["aciklama"])
        self.assertIn("borcu var", r["aciklama"])

    def test_diger_tesvikler(self):
        p = self.p
        self.assertEqual(main.t_81i(kisi(), isyeri(), p["5510_81_i"])["aciklama"], "2 puan (imalat dışı işyeri)")
        self.assertEqual(main.t_81i(kisi(), isyeri(imalat=True), p["5510_81_i"])["aciklama"], "5 puan (imalat işyeri)")
        self.assertEqual(main.t_81i(kisi(), isyeri(kayit_disi="Evet"), p["5510_81_i"])["durum"], main.DEGIL)
        self.assertEqual(main.t_81i(kisi(), isyeri(borc=""), p["5510_81_i"])["durum"], main.EKSIK)
        self.assertEqual(main.t_engelli(kisi(engelli=True), isyeri(), p["4857_30_engelli"])["durum"], main.UYGUN)
        self.assertEqual(main.t_engelli(kisi(engelli=True), isyeri(surede="Hayır"), p["4857_30_engelli"])["durum"], main.DEGIL)
        self.assertEqual(main.t_md50(kisi(odenek=True, son_isyeri=False), isyeri(), p["4447_50_issizlik_odenegi"])["durum"], main.UYGUN)
        self.assertIn("son çalıştığı", main.t_md50(kisi(odenek=True, son_isyeri=True), isyeri(), p["4447_50_issizlik_odenegi"])["aciklama"])
        self.assertEqual(main.t_md50(kisi(odenek=True), isyeri(), p["4447_50_issizlik_odenegi"])["durum"], main.EKSIK)

    def test_ornek_ve_rapor(self):
        with tempfile.TemporaryDirectory() as tmp:
            r = main.calistir(CAL, ISY, PAR, Path(tmp) / "t.xlsx")
            d = {s["sicil"]: {k: v["durum"] for k, v in s["sonuclar"].items()} for s in r["satirlar"]}
            self.assertEqual([d[x]["4447_gecici_10"] for x in sorted(d)], ["Uygun", "Uygun", "Uygun", "Uygun değil", "Uygun değil", "Uygun değil", "Uygun değil"])
            self.assertEqual((d["P004"]["4857_30_engelli"], d["P003"]["4447_50_issizlik_odenegi"], d["P007"]["4447_50_issizlik_odenegi"]),
                             ("Uygun", "Uygun", "Bilgi eksik"))
            ws = load_workbook(Path(tmp) / "t.xlsx")["Uygunluk"]
            self.assertEqual((ws["A2"].value, ws["F2"].value), ("P001", "Uygun: 30 ay · 01.09.2026-28.02.2029"))
            self.assertEqual(load_workbook(Path(tmp) / "t.xlsx")["Özet"]["C3"].value, 3)

    def test_cli_ve_hatalar(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(main.main(["--calisanlar", str(Path(tmp) / "yok.csv")]), 1)
            eksik = Path(tmp) / "p.json"
            eksik.write_text(json.dumps({"5510_81_i": {}}), encoding="utf-8")
            self.assertEqual(main.main(["--parametre", str(eksik)]), 1)
            bozuk = Path(tmp) / "c.csv"
            bozuk.write_text("Ad Soyad;Doğum Tarihi;Cinsiyet;İşe Giriş Tarihi\nX;bilinmiyor;Kadın;01.10.2026\n", encoding="utf-8")
            self.assertEqual(main.main(["--calisanlar", str(bozuk)]), 1)


if __name__ == "__main__":
    unittest.main()
