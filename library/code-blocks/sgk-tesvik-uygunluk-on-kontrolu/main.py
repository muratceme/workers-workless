"""
SGK Teşvik Uygunluk Ön Kontrolü — Gerisi Bende kod bloğu
Bordro ve SGK İşlemleri › Bordro Elemanı
Çalışan ve işyeri bilgilerine göre yaygın SGK prim teşviklerinin uygunluk koşullarını ön kontrolden geçirir.
Tarih ve oranlar tesvik_parametreleri.json dosyasındadır; mevzuat değiştiğinde dosya güncellenir.

  Ortak işyeri şartları (81/ı, geçici 10): özel sektör · bildirgeler yasal sürede · prim borcu yok ya da yapılandırılmış · kayıt dışı tespiti yok
  5510 md. 81/1-(ı)  = her sigortalı; MYÖ işveren hissesinde imalat / diğer puan indirimi (parametre)
  4447 geçici 10     = işe giriş [ilk_ise_alim, son_ise_alim] · 18 yaşını doldurmuş · işe girişten önceki 6 ayda kayıtlı değil ·
                       ortalama sigortalı sayısına ilave. Süre: kadın ya da 29 yaşını doldurmamış erkek → MYB 48, mesleki eğitim /
                       İŞKUR kursu 36, belgesiz 24 ay (+6 İŞKUR kayıtlı); 29 yaşını doldurmuş erkek → belgeli 24 (+6 İŞKUR), belgesiz
                       ama İŞKUR kayıtlı 6 ay
  4857 md. 30        = engelli sigortalı; bildirge yasal sürede (PEK alt sınırı üzerinden işveren hissesinin tamamı)
  4447 md. 50        = işe giriş ≥ ilk_ise_alim · işe girişte işsizlik ödeneği hakkı · ödenek öncesi son işyeri bu işyeri değil ·
                       ortalamaya ilave · bildirge yasal sürede (süre: kalan ödenek süresi)
  Sonuç: Uygun / Uygun değil (nedenleriyle) / Bilgi eksik. Aynı sigortalı ve aynı prim günü için birden fazla teşvikin
  birlikte uygulanıp uygulanamayacağı ayrıca değerlendirilir.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
KIRMIZI, SARI, YESIL = "FDE2E1", "FFF4CE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
UYGUN, DEGIL, EKSIK = "Uygun", "Uygun değil", "Bilgi eksik"

SUTUN = {
    "sicil": ("Sicil No", "Personel No"), "ad": ("Ad Soyad", "Adı Soyadı"), "dogum": ("Doğum Tarihi",), "cinsiyet": ("Cinsiyet",),
    "giris": ("İşe Giriş Tarihi", "Giriş Tarihi"), "son6": ("Son 6 Ayda Sigortalı mı", "Son 6 Ay Sigortalı"), "iskur": ("İŞKUR Kaydı", "İşkur Kayıtlı"),
    "belge": ("Mesleki Belge", "Belge"), "engelli": ("Engelli",), "odenek": ("İşsizlik Ödeneği", "İşsizlik Ödeneği Alıyor"),
    "son_isyeri": ("Ödenek Öncesi Son İşyeri Burası mı", "Son İşyeri Burası"), "ortalama": ("Önceki 6 Ay Ortalama Sigortalı", "Ortalama Sigortalı"),
    "sayi": ("Giriş Ayı Sigortalı Sayısı", "Sigortalı Sayısı"),
}
ISYERI = {"tarih": "Kontrol Tarihi", "sektor": "Sektör", "ozel": "Özel Sektör", "surede": "Bildirgeler Yasal Sürede", "borc": "Prim Borcu",
          "kayit_disi": "Kayıt Dışı Tespiti"}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def evet(x) -> bool | None:
    k = katla(x)
    if k in {"evet", "e", "var", "1", "true", "x"}:
        return True
    if k in {"hayir", "h", "yok", "0", "false"}:
        return False
    return None


def sayi(x) -> float | None:
    try:
        return float(str(x).replace(",", ".")) if str(x or "").strip() else None
    except ValueError:
        return None


def yas(dogum: date, gun: date) -> int:
    return gun.year - dogum.year - ((gun.month, gun.day) < (dogum.month, dogum.day))


def ay_ekle(g: date, ay: int) -> date:
    y, m = divmod(g.month - 1 + ay, 12)
    yil, ay_ = g.year + y, m + 1
    gun = min(g.day, [31, 29 if yil % 4 == 0 and (yil % 100 or yil % 400 == 0) else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31][ay_ - 1])
    return date(yil, ay_, gun)


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def calisanlari_oku(yol: Path) -> list[dict]:
    s = tablo_oku(yol)
    if not s:
        raise ValueError(f"{yol.name}: tablo boş")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in SUTUN.items()}
    eksik = [SUTUN[a][0] for a in ("ad", "dogum", "cinsiyet", "giris") if idx[a] is None]
    if eksik:
        raise ValueError(f"{yol.name}: gerekli sütun yok: {', '.join(eksik)}")
    out = []
    for r in s[1:]:
        v = {a: (r[i] if i is not None and i < len(r) else "") for a, i in idx.items()}
        c = {"sicil": str(v["sicil"] or ""), "ad": str(v["ad"] or ""), "dogum": tarih(v["dogum"]), "giris": tarih(v["giris"]),
             "kadin": katla(v["cinsiyet"]).startswith("k"), "cinsiyet_bilinir": katla(v["cinsiyet"])[:1] in {"k", "e"},
             "son6": evet(v["son6"]), "iskur": evet(v["iskur"]), "engelli": evet(v["engelli"]), "odenek": evet(v["odenek"]),
             "son_isyeri": evet(v["son_isyeri"]), "ortalama": sayi(v["ortalama"]), "sayi": sayi(v["sayi"]), "belge": belge_turu(v["belge"])}
        if not c["dogum"] or not c["giris"]:
            raise ValueError(f"{yol.name}: {c['ad'] or c['sicil']}: doğum / işe giriş tarihi okunamadı")
        out.append(c)
    return out


def belge_turu(x) -> str:
    k = katla(x)
    if not k or k in {"yok", "hayir", "0"}:
        return "yok"
    if "yeterli" in k or "myb" in k:
        return "myb"
    return "mesleki_egitim"  # mesleki ve teknik eğitim mezuniyeti ya da İŞKUR işgücü yetiştirme kursu


def isyeri_oku(yol: Path) -> dict:
    i = {a: "" for a in ISYERI}
    for s in yol.read_text(encoding="utf-8-sig").splitlines():
        a, _, d = s.partition(":")
        k = next((k for k, v in ISYERI.items() if katla(v) == katla(a)), None)
        if k:
            i[k] = d.strip()
    i["tarih_v"] = tarih(i["tarih"]) or date.today()
    i["imalat"] = katla(i["sektor"]).startswith("imalat")
    return i


def ortak_isyeri(i: dict) -> tuple[list[str], list[str]]:
    """(engel nedenleri, eksik bilgiler)"""
    engel, eksik = [], []
    for alan, beklenen, mesaj in (("ozel", True, "özel sektör işyeri değil"), ("surede", True, "bildirgeler / primler yasal sürede değil"),
                                  ("kayit_disi", False, "kayıt dışı sigortalı / sahte bildirim tespiti var")):
        v = evet(i[alan])
        if v is None:
            eksik.append(ISYERI[alan])
        elif v != beklenen:
            engel.append(mesaj)
    b = katla(i["borc"])
    if not b:
        eksik.append(ISYERI["borc"])
    elif b not in {"yok", "hayir", "yapilandirilmis", "taksitlendirilmis", "yapilandirma"}:
        engel.append("vadesi geçmiş prim / idari para cezası borcu var (yapılandırılmamış)")
    return engel, eksik


def ilave_mi(c: dict) -> bool | None:
    if c["ortalama"] is None or c["sayi"] is None:
        return None
    return c["sayi"] > c["ortalama"]


def sonuc(engel: list[str], eksik: list[str], uygun_aciklama: str) -> dict:
    if engel:
        return {"durum": DEGIL, "aciklama": "; ".join(engel)}
    if eksik:
        return {"durum": EKSIK, "aciklama": "Eksik: " + ", ".join(eksik)}
    return {"durum": UYGUN, "aciklama": uygun_aciklama}


def t_81i(c: dict, i: dict, p: dict) -> dict:
    engel, eksik = ortak_isyeri(i)
    puan = p["puan_imalat"] if i["imalat"] else p["puan_diger"]
    return sonuc(engel, eksik, f"{puan} puan ({'imalat' if i['imalat'] else 'imalat dışı'} işyeri)")


def t_gecici10(c: dict, i: dict, p: dict) -> dict:
    engel, eksik = ortak_isyeri(i)
    if not tarih(p["ilk_ise_alim"]) <= c["giris"] <= tarih(p["son_ise_alim"]):
        engel.append(f"işe giriş {p['ilk_ise_alim']}-{p['son_ise_alim']} aralığında değil")
    y = yas(c["dogum"], c["giris"])
    if y < 18:
        engel.append(f"işe girişte {y} yaşında (18 yaşını doldurmamış)")
    if c["son6"] is None:
        eksik.append("son 6 ayda sigortalılık")
    elif c["son6"]:
        engel.append("işe girişten önceki 6 ayda sigortalı olarak kayıtlı")
    il = ilave_mi(c)
    if il is None:
        eksik.append("ortalama / giriş ayı sigortalı sayısı")
    elif not il:
        engel.append(f"ortalama sigortalı sayısına ilave değil ({c['sayi']:.0f} ≤ {c['ortalama']:.0f})")
    if not c["cinsiyet_bilinir"]:
        eksik.append("cinsiyet")
    sure = 0
    if y >= 18 and c["cinsiyet_bilinir"]:
        iskur = bool(c["iskur"])
        if c["kadin"] or y < 29:
            sure = p["sure_genc_kadin"][c["belge"]] + (p["iskur_ilave_ay"] if iskur else 0)
        elif c["belge"] != "yok":
            sure = p["sure_29_ustu_erkek"]["belgeli"] + (p["iskur_ilave_ay"] if iskur else 0)
        elif iskur:
            sure = p["sure_29_ustu_erkek"]["iskur_belgesiz"]
        elif c["iskur"] is None:
            eksik.append("İŞKUR kaydı")
        else:
            engel.append("29 yaşını doldurmuş erkek; mesleki belge ve İŞKUR kaydı yok")
    r = sonuc(engel, eksik, "")
    if r["durum"] == UYGUN:
        bitis = ay_ekle(c["giris"], sure) - timedelta(days=1)
        r["aciklama"] = f"{sure} ay · {c['giris']:%d.%m.%Y}-{bitis:%d.%m.%Y}"
        r["bitis"] = bitis
        if bitis < i["tarih_v"]:
            r = {"durum": DEGIL, "aciklama": f"süre {bitis:%d.%m.%Y} tarihinde doldu"}
    return r


def t_engelli(c: dict, i: dict, p: dict) -> dict:
    engel, eksik = [], []
    if c["engelli"] is None:
        eksik.append("engellilik durumu")
    elif not c["engelli"]:
        engel.append("engelli sigortalı değil")
    v = evet(i["surede"])
    if v is None:
        eksik.append(ISYERI["surede"])
    elif not v:
        engel.append("bildirgeler / primler yasal sürede değil")
    return sonuc(engel, eksik, "PEK alt sınırı üzerinden işveren hissesinin tamamı (engellilik raporu dosyada olmalı)")


def t_md50(c: dict, i: dict, p: dict) -> dict:
    engel, eksik = [], []
    if c["odenek"] is None:
        eksik.append("işsizlik ödeneği durumu")
    elif not c["odenek"]:
        engel.append("işe girişte işsizlik ödeneği hakkı yok")
    if c["giris"] < tarih(p["ilk_ise_alim"]):
        engel.append(f"işe giriş {p['ilk_ise_alim']} öncesi")
    if c["odenek"]:
        if c["son_isyeri"] is None:
            eksik.append("ödenek öncesi son işyeri")
        elif c["son_isyeri"]:
            engel.append("ödenekten önce son çalıştığı işyeri burası")
    il = ilave_mi(c)
    if il is None:
        eksik.append("ortalama / giriş ayı sigortalı sayısı")
    elif not il:
        engel.append("ortalama sigortalı sayısına ilave değil")
    v = evet(i["surede"])
    if v is None:
        eksik.append(ISYERI["surede"])
    elif not v:
        engel.append("bildirgeler / primler yasal sürede değil")
    return sonuc(engel, eksik, "kalan işsizlik ödeneği süresince (süreyi İŞKUR kaydından kontrol edin)")


TESVIKLER = [("5510_81_i", t_81i), ("4447_gecici_10", t_gecici10), ("4857_30_engelli", t_engelli), ("4447_50_issizlik_odenegi", t_md50)]


def parametre_oku(yol: Path) -> dict:
    try:
        p = json.loads(yol.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as h:
        raise ValueError(f"{yol.name}: parametre dosyası okunamadı ({h})")
    eksik = [k for k, _ in TESVIKLER if k not in p]
    if eksik:
        raise ValueError(f"{yol.name}: parametre eksik: {', '.join(eksik)}")
    return p


def analiz(calisanlar: list[dict], i: dict, p: dict) -> list[dict]:
    out = []
    for c in calisanlar:
        out.append(c | {"yas": yas(c["dogum"], c["giris"]), "sonuclar": {k: f(c, i, p[k]) for k, f in TESVIKLER}})
    return out


def rapor_yaz(cikti: Path, satirlar: list[dict], i: dict, p: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Uygunluk"
    bas = ["Sicil No", "Ad Soyad", "İşe Giriş", "Girişte Yaş"] + [p[k]["dayanak"] for k, _ in TESVIKLER]
    ws.append(bas)
    for h in ws[1]:
        h.fill, h.font, h.alignment = BASLIK_DOLGU, BASLIK_YAZI, UST
    for j, w in enumerate([10, 22, 12, 8] + [42] * len(TESVIKLER), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for s in satirlar:
        ws.append([s["sicil"], s["ad"], s["giris"], s["yas"]] + [f"{r['durum']}: {r['aciklama']}" if r["aciklama"] else r["durum"] for r in s["sonuclar"].values()])
        ws.cell(ws.max_row, 3).number_format = "DD.MM.YYYY"
        for j, r in enumerate(s["sonuclar"].values(), 5):
            ws.cell(ws.max_row, j).fill = PatternFill("solid", fgColor={UYGUN: YESIL, DEGIL: KIRMIZI, EKSIK: SARI}[r["durum"]])
    ws.freeze_panes = "C2"
    for satir in ws.iter_rows(min_row=2):
        for h in satir:
            h.alignment = UST
    ws.append([])
    ws.append([f"Kontrol tarihi: {i['tarih_v']:%d.%m.%Y} · Parametrelerin son teyidi: {p.get('son_teyit', '?')} · Ön kontroldür; teşvik tanımlaması "
               "ve bildirge öncesi SGK e-Bildirge / teşvik sorgulama ekranlarıyla ve güncel mevzuatla doğrulayın. Aynı sigortalı ve aynı prim günü "
               "için birden fazla teşvikin birlikte uygulanıp uygulanamayacağını ayrıca değerlendirin."])
    o = wb.create_sheet("Özet")
    o.append(["Teşvik", "Dayanak", "Uygun", "Uygun değil", "Bilgi eksik", "Not"])
    for h in o[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for k, _ in TESVIKLER:
        d = [s["sonuclar"][k]["durum"] for s in satirlar]
        o.append([p[k]["ad"], p[k]["dayanak"], d.count(UYGUN), d.count(DEGIL), d.count(EKSIK), p[k].get("not", "")])
    for j, w in enumerate((40, 34, 8, 12, 12, 70), 1):
        o.column_dimensions[get_column_letter(j)].width = w
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(calisan_yol: Path, isyeri_yol: Path, parametre_yol: Path, cikti: Path) -> dict:
    p = parametre_oku(parametre_yol)
    i = isyeri_oku(isyeri_yol)
    satirlar = analiz(calisanlari_oku(calisan_yol), i, p)
    rapor_yaz(cikti, satirlar, i, p)
    return {"satirlar": satirlar, "isyeri": i, "parametre": p}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Çalışan ve işyeri bilgilerine göre SGK prim teşviklerinin uygunluk koşullarını ön kontrolden geçirir.")
    p.add_argument("--calisanlar", type=Path, default=ORNEK / "calisanlar.csv",
                   help="Sicil No, Ad Soyad, Doğum Tarihi, Cinsiyet, İşe Giriş Tarihi, Son 6 Ayda Sigortalı mı, İŞKUR Kaydı, Mesleki Belge, Engelli, "
                        "İşsizlik Ödeneği, Ödenek Öncesi Son İşyeri Burası mı, Önceki 6 Ay Ortalama Sigortalı, Giriş Ayı Sigortalı Sayısı")
    p.add_argument("--isyeri", type=Path, default=ORNEK / "isyeri.txt",
                   help="Kontrol Tarihi, Sektör (İmalat / Diğer), Özel Sektör, Bildirgeler Yasal Sürede, Prim Borcu (Yok / Yapılandırılmış / Var), Kayıt Dışı Tespiti")
    p.add_argument("--parametre", type=Path, default=ORNEK / "tesvik_parametreleri.json", help="Teşvik tarih ve oranları (JSON); mevzuat değişince güncelleyin")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "tesvik_uygunluk.xlsx")
    a = p.parse_args(argv)
    for yol in (a.calisanlar, a.isyeri, a.parametre):
        if not yol.exists():
            print(f"[X] Dosya bulunamadı: {yol}")
            return 1
    try:
        r = calistir(a.calisanlar, a.isyeri, a.parametre, a.cikti)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(r['satirlar'])} çalışan · parametrelerin son teyidi {r['parametre'].get('son_teyit', '?')}")
    for k, _ in TESVIKLER:
        d = [s["sonuclar"][k]["durum"] for s in r["satirlar"]]
        print(f"    {r['parametre'][k]['dayanak']}: uygun {d.count(UYGUN)} · uygun değil {d.count(DEGIL)} · bilgi eksik {d.count(EKSIK)}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
