import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 9)


def kural(tip, urun, belge, ay=None):
    return {"tip": main.katla(tip), "urun": main.katla(urun), "belge": belge, "ay": ay, "aciklama": ""}


KURALLAR = [kural("Gerçek Kişi", "*", "Kimlik Belgesi"), kural("Tüzel Kişi", "*", "Faaliyet Belgesi", 6), kural("*", "Kredi", "Mali Durum Belgesi", 12),
            kural("Tüzel Kişi", "Kredi", "Mali Durum Belgesi", 6), kural("Tüzel Kişi", "*", "İmza Sirküleri")]


def musteri(no="M1", tip="Gerçek Kişi", urunler=("Vadesiz",), risk="Düşük"):
    return {"no": no, "ad": "Test", "tip": tip, "urunler": list(urunler), "risk": risk, "sube": "Merkez", "py": "PY"}


def belge(ad, t=None, bitis=None):
    return {"belge": ad, "tarih": t, "bitis": bitis}


def durumlar(r):
    return {(s["m"]["no"], s["belge"]): s["durum"] for s in r["satirlar"]}


class KuralTesti(unittest.TestCase):
    def test_tip_ve_urun_eslesmesi(self):
        g = main.gerekenler(musteri(tip="TÜZEL KİŞİ", urunler=("Vadesiz", "kredi")), KURALLAR)
        self.assertEqual(sorted(x["belge"] for x in g.values()), ["Faaliyet Belgesi", "Mali Durum Belgesi", "İmza Sirküleri"])
        self.assertEqual(g["mali durum belgesi"]["ay"], 6)                          # iki kuraldan kısa olan
        self.assertEqual(list(main.gerekenler(musteri(), KURALLAR)), ["kimlik belgesi"])

    def test_eksik_ve_sureli(self):
        m = musteri(tip="Tüzel Kişi")
        d = {"M1": [belge("FAALİYET BELGESİ", date(2026, 4, 9))]}                   # + 6 ay = 09.10.2026 → bugün bitiyor
        r = main.kontrol([m], d, KURALLAR, BUGUN)
        self.assertEqual(durumlar(r), {("M1", "Faaliyet Belgesi"): "Yaklaşan", ("M1", "İmza Sirküleri"): "Eksik"})
        d = {"M1": [belge("Faaliyet Belgesi", date(2026, 4, 8)), belge("İmza Sirküleri")]}
        self.assertEqual(durumlar(main.kontrol([m], d, KURALLAR, BUGUN)), {("M1", "Faaliyet Belgesi"): "Süresi dolmuş", ("M1", "İmza Sirküleri"): "Tamam"})
        d = {"M1": [belge("Faaliyet Belgesi"), belge("İmza Sirküleri", bitis=date(2026, 11, 20))]}
        self.assertEqual(durumlar(main.kontrol([m], d, KURALLAR, BUGUN)), {("M1", "Faaliyet Belgesi"): "Tarih yok", ("M1", "İmza Sirküleri"): "Tamam"})
        self.assertEqual(durumlar(main.kontrol([m], d, KURALLAR, BUGUN, 60))[("M1", "İmza Sirküleri")], "Yaklaşan")

    def test_bitis_sutunu_ve_en_yeni_kayit(self):
        m = musteri()
        d = {"M1": [belge("Kimlik Belgesi", date(2016, 1, 1), date(2026, 1, 1)), belge("Kimlik Belgesi", date(2026, 2, 1), date(2036, 2, 1))]}
        r = main.kontrol([m], d, KURALLAR, BUGUN)
        self.assertEqual((r["satirlar"][0]["durum"], r["satirlar"][0]["bitis"], r["satirlar"][0]["tarih"]), ("Tamam", date(2036, 2, 1), date(2026, 2, 1)))
        d = {"M1": [belge("Kimlik Belgesi", date(2016, 1, 1), date(2026, 1, 1))]}
        self.assertEqual(main.kontrol([m], d, KURALLAR, BUGUN)["satirlar"][0]["kalan"], -281)

    def test_ay_ekle(self):
        self.assertEqual(main.ay_ekle(date(2025, 8, 31), 6), date(2026, 2, 28))
        self.assertEqual(main.ay_ekle(date(2024, 2, 29), 12), date(2025, 2, 28))
        self.assertEqual(main.ay_ekle(date(2026, 11, 15), 3), date(2027, 2, 15))

    def test_siralama_ve_ozet(self):
        ms = [musteri("M1"), musteri("M2", risk="Yüksek"), musteri("M3", tip="Vakıf")]
        r = main.kontrol(ms, {"M1": [belge("Kimlik Belgesi")], "M3": [belge("Pasaport")]}, KURALLAR, BUGUN)
        self.assertEqual(r["satirlar"][0]["m"]["no"], "M2")                         # yüksek riskli ve eksik önce
        self.assertEqual({o["m"]["no"]: (o["durum"], o["oran"]) for o in r["ozet"]}, {"M1": ("Tam", 100), "M2": ("Eksik var", 0), "M3": ("Kural yok", None)})
        self.assertEqual(sorted(u["tur"] for u in r["uyarilar"]), ["Kuralda olmayan belge", "Müşteri tipi için kural yok", "Ürün için özel kural yok"])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "k.xlsx"
            r = main.calistir(ORNEK / "musteriler.csv", ORNEK / "belgeler.csv", ORNEK / "kurallar.csv", cikti, BUGUN)
            say = {d: sum(s["durum"] == d for s in r["satirlar"]) for d in main.DURUM_SIRA}
            self.assertEqual((len(r["ozet"]), say["Eksik"], say["Süresi dolmuş"], say["Yaklaşan"]), (19, 6, 4, 3))
            self.assertEqual(r["satirlar"][0]["m"]["risk"], "Yüksek")
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Eksik Listesi", "Müşteri Özeti", "Belge Özeti", "Şube Özeti", "Uyarılar"])
            self.assertEqual(wb["Eksik Listesi"].max_row, 1 + 6 + 4 + 3)

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            m, b, k = (Path(tmp) / x for x in ("m.csv", "b.csv", "k.csv"))
            m.write_text("Müşteri No;Unvan;Müşteri Tipi;Ürünler\nM1;A;Gerçek Kişi;Vadesiz, Kredi\nM1;A;Gerçek Kişi;Vadesiz\n", encoding="utf-8")
            b.write_text("Müşteri No;Belge;Belge Tarihi\nM1;Kimlik;dün\nM9;Kimlik;01.01.2026\n", encoding="utf-8")
            k.write_text("Müşteri Tipi;Ürün;Belge;Geçerlilik (ay)\n*;*;Kimlik;on iki\n", encoding="utf-8")
            ms, bs, ks, uy = main.oku(m, b, k)
            self.assertEqual((ms[0]["urunler"], len(ms), ks[0]["ay"]), (["Vadesiz", "Kredi"], 1, None))
            self.assertEqual(sorted(u["tur"] for u in uy), ["Kural geçerlilik süresi okunamadı", "Mükerrer müşteri", "Müşteri listesinde yok", "Tarih okunamadı"])
            arg = ["--musteriler", str(m), "--belgeler", str(b), "--kurallar", str(k), "--cikti", str(Path(tmp) / "x.xlsx")]
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(arg + ["--bugun", "bugün"]), 2)
            self.assertEqual(main.main(["--kurallar", str(Path(tmp) / "yok.csv")]), 1)
            k.write_text("Müşteri Tipi;Ürün;Belge\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)


if __name__ == "__main__":
    unittest.main()
