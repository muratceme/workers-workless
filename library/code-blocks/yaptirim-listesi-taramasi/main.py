"""
Yaptırım Listesi Taraması — Workers / Workless kod bloğu
Bankacılık › Kurumsal Uyum › Uyum Uzmanı

Müşteri ve karşı taraf listesini yaptırım listeleriyle bulanık eşleştirir:
  - Liste kaynakları: BM Güvenlik Konseyi Konsolide Listesi (resmî XML), ya da kendi derlediğiniz CSV/Excel
    (Resmî Gazete'de yayımlanan malvarlığı dondurma kararları, OFAC, AB vb.): Ad, Diğer Adlar, Doğum Tarihi, Uyruk, Liste, Referans.
  - Ad normalleştirme: Türkçe karakter ve büyük/küçük harf farkı, aksanlar, noktalama, şirket ekleri (A.Ş., LTD. ŞTİ.,
    LLC…), yaygın transliterasyon farkları (Mohammed/Muhammed/Mehmet, Hussein/Hüseyin, Yousef/Yusuf, Ahmed/Ahmet…).
  - Benzerlik: kelime bazında bulanık eşleşme (sıra bağımsız) ve takma adların her biri; doğum tarihi/yılı ve uyruk
    eşleşmesi puanı artırır, çelişmesi düşürür.
  - Sonuç seviyeleri: Güçlü eşleşme · Olası eşleşme · Zayıf benzerlik; analist karar sütunları.
İnternete bağlanmaz: listeleri kendiniz indirip verirsiniz.

Kullanım:
    python main.py                                         # kurgusal örnek verilerle
    python main.py --musteriler musteriler.xlsx --liste consolidated.xml --liste resmi_gazete_kararlari.xlsx
    python main.py --musteriler musteriler.xlsx --liste liste.csv --esik 80
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
import unicodedata
from collections import defaultdict
from difflib import SequenceMatcher
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent

SIRKET_EKLERI = {"as", "a s", "anonim", "sirketi", "sti", "ltd", "limited", "llc", "inc", "co", "corp", "corporation", "company",
                 "gmbh", "sa", "sarl", "plc", "san", "tic", "ve", "and", "the", "holding", "group", "grup", "fze", "llp"}
UNVAN = {"mr", "mrs", "ms", "dr", "prof", "haci", "hacı", "sheikh", "seyh", "molla", "mullah", "imam", "al", "el", "bin", "ibn",
         "bint", "abu", "ebu", "oglu"}
# Yaygın adların transliterasyon biçimleri → ortak kök (yalnız tam kelime eşleşmesinde uygulanır)
ESDEGER = {}
for kok, bicimler in {
    "muhammed": "mohammed mohammad mohamed mohamad muhammad muhamed mehmet mehmed mohd muhammet",
    "ahmet": "ahmed ahmad achmed",
    "huseyin": "hussein husein husain hussain husayn hossein",
    "hasan": "hassan hasen",
    "yusuf": "yousef yousif youssef yusef josef",
    "omer": "omar umar",
    "osman": "othman uthman usman",
    "abdullah": "abdallah abdulla",
    "mustafa": "mostafa moustafa mustapha",
    "ibrahim": "ebrahim ibrahem",
    "ismail": "esmail ismael",
    "suleyman": "sulaiman suleiman soliman",
    "halil": "khalil khaleel",
    "halid": "khalid khaled halit",
    "abdurrahman": "abdulrahman abdelrahman abdul rahman",
    "yunus": "younes younis yunis",
    "ali": "aly",
}.items():
    for b in bicimler.split():
        ESDEGER[b] = kok
    ESDEGER[kok] = kok


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def katla(s) -> str:
    """Türkçe ve diğer aksanlı harfleri ASCII'ye indirger; noktalama ve fazla boşluk atılır."""
    s = kucuk(s).replace("ı", "i")
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = s.replace("'", "").replace("’", "").replace("`", "")
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def ses_normal(kelime: str) -> str:
    """Hafif ses normalleştirmesi: çift harf, ph→f, w→v, q→k, kh→h, th→t, ou→u, ee→i."""
    k = ESDEGER.get(kelime, kelime)
    for a, b in (("kh", "h"), ("ph", "f"), ("th", "t"), ("ou", "u"), ("ee", "i"), ("oo", "u"), ("w", "v"), ("q", "k"), ("x", "ks")):
        k = k.replace(a, b)
    return re.sub(r"(.)\1+", r"\1", k)


def ad_kelimeleri(ad: str, sirket: bool = False) -> list[str]:
    kel = [k for k in katla(ad).split() if k]
    atilacak = UNVAN | (SIRKET_EKLERI if sirket else set())
    kel = [k for k in kel if k not in atilacak and not (sirket and len(k) == 1)]
    return [ses_normal(k) for k in kel]


def benzerlik(a: list[str], b: list[str]) -> float:
    """Sıra bağımsız kelime eşleşmesi (0–100): her kelime diğer addaki en benzer kelimeyle eşlenir (≥ 0,80)."""
    if not a or not b:
        return 0.0
    kalan = list(b)
    toplam = 0.0
    for k in sorted(a, key=len, reverse=True):
        en_iyi, j_iyi = 0.0, None
        for j, m in enumerate(kalan):
            r = 1.0 if k == m else SequenceMatcher(None, k, m).ratio()
            if len(k) == 1 or len(m) == 1:                     # baş harf: "M." ↔ "Mehmet"
                r = 0.85 if k[0] == m[0] else 0.0
            if r > en_iyi:
                en_iyi, j_iyi = r, j
        if en_iyi >= 0.80 and j_iyi is not None:
            toplam += en_iyi
            kalan.pop(j_iyi)
    return 100 * 2 * toplam / (len(a) + len(b))


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254", "latin-1"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def _bul(baslik, *adlar):
    b = [katla(x) for x in baslik]
    return next((b.index(katla(a)) for a in adlar if katla(a) in b), None)


def _al(r, i):
    return r[i] if i is not None and i < len(r) else None


def iso(x) -> str:
    """Tam doğum tarihini YYYY-AA-GG biçimine çevirir; yalnız yıl veya okunamayan değer için ''."""
    t = str(x or "").strip()
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", t)
    if m:
        return f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    m = re.match(r"^(\d{1,2})[./](\d{1,2})[./](\d{4})", t)
    if m:
        return f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
    return ""


def yil(x) -> str:
    m = re.search(r"(19|20)\d{2}", str(x or ""))
    return m.group(0) if m else ""


# ----------------------------------------------------------------------------
# Liste okuma
# ----------------------------------------------------------------------------

def bm_xml_oku(yol: Path) -> list[dict]:
    """BM Güvenlik Konseyi Konsolide Listesi: CONSOLIDATED_LIST > INDIVIDUALS > INDIVIDUAL, ENTITIES > ENTITY."""
    kok = ElementTree.parse(yol).getroot()
    kayitlar = []
    for tur, etiket in (("Kişi", "INDIVIDUAL"), ("Kuruluş", "ENTITY")):
        for e in kok.iter(etiket):
            adlar = [(e.findtext(t) or "").strip() for t in ("FIRST_NAME", "SECOND_NAME", "THIRD_NAME", "FOURTH_NAME")]
            ad = " ".join(a for a in adlar if a)
            diger = [(x.text or "").strip() for x in e.iter("ALIAS_NAME") if (x.text or "").strip()]
            dogum = sorted({(d.findtext("DATE") or d.findtext("YEAR") or "").strip()[:10]
                            for d in e.iter(f"{etiket}_DATE_OF_BIRTH")} - {""})
            uyruk = [(v.text or "").strip() for n in e.iter("NATIONALITY") for v in n.iter("VALUE") if (v.text or "").strip()]
            kayitlar.append({"ad": ad, "diger": diger, "dogum": dogum, "uyruk": uyruk, "tur": tur,
                             "liste": "BM Konsolide Liste", "ref": (e.findtext("REFERENCE_NUMBER") or "").strip()})
    return kayitlar


def liste_oku(yol: Path) -> list[dict]:
    if yol.suffix.lower() == ".xml":
        return bm_xml_oku(yol)
    s = tablo_oku(yol)
    b = s[0]
    i = {k: _bul(b, *v) for k, v in {
        "ad": ("ad", "ad soyad unvan", "ad soyad", "unvan", "name", "isim"), "diger": ("diğer adlar", "takma adlar", "alias", "aliases", "aka"),
        "dogum": ("doğum tarihi", "doğum yılı", "dob", "date of birth"), "uyruk": ("uyruk", "nationality", "ülke"),
        "tur": ("tür", "tip", "type"), "liste": ("liste", "kaynak", "program", "list"), "ref": ("referans", "karar", "reference", "ref")}.items()}
    if i["ad"] is None:
        raise SystemExit(f"{yol.name}: 'Ad' sütunu bulunamadı. Başlıklar: {b}")
    kayitlar = []
    for r in s[1:]:
        ad = str(_al(r, i["ad"]) or "").strip()
        if not ad:
            continue
        tur = katla(_al(r, i["tur"]))
        kayitlar.append({"ad": ad, "diger": [x.strip() for x in re.split(r"[;|]", str(_al(r, i["diger"]) or "")) if x.strip()],
                         "dogum": [str(_al(r, i["dogum"])).strip()] if _al(r, i["dogum"]) else [],
                         "uyruk": [x.strip() for x in re.split(r"[;,|]", str(_al(r, i["uyruk"]) or "")) if x.strip()],
                         "tur": "Kuruluş" if tur.startswith(("kurulus", "entity", "sirket", "tuzel")) else "Kişi",
                         "liste": str(_al(r, i["liste"]) or yol.stem), "ref": str(_al(r, i["ref"]) or "")})
    return kayitlar


def musteriler_oku(yol: Path) -> list[dict]:
    s = tablo_oku(yol)
    b = s[0]
    i = {k: _bul(b, *v) for k, v in {
        "no": ("müşteri no", "no", "id", "cari kod"), "ad": ("ad soyad unvan", "ad soyad", "unvan", "müşteri", "ad", "isim"),
        "dogum": ("doğum tarihi", "kuruluş tarihi", "doğum yılı"), "uyruk": ("uyruk", "ülke"), "tur": ("tip", "tür", "müşteri tipi"),
        "rol": ("rol", "ilişki", "taraf")}.items()}
    if i["ad"] is None:
        raise SystemExit(f"{yol.name}: 'Ad Soyad / Unvan' sütunu bulunamadı. Başlıklar: {b}")
    return [{"no": str(_al(r, i["no"]) or n), "ad": str(_al(r, i["ad"])).strip(), "dogum": str(_al(r, i["dogum"]) or "").strip(),
             "uyruk": str(_al(r, i["uyruk"]) or "").strip(), "rol": str(_al(r, i["rol"]) or "").strip(),
             "tur": "Kuruluş" if katla(_al(r, i["tur"])).startswith(("kurum", "tuzel", "sirket", "kurulus")) else "Kişi"}
            for n, r in enumerate(s[1:], 1) if _al(r, i["ad"])]


# ----------------------------------------------------------------------------
# Eşleştirme
# ----------------------------------------------------------------------------

def tara(musteriler: list[dict], liste: list[dict], esik: float = 85.0, zayif: float = 75.0) -> list[dict]:
    # Ön indeks: kelimenin ilk iki harfi → kayıtlar (karşılaştırma sayısını azaltır)
    hazir = []
    indeks: dict[str, set[int]] = defaultdict(set)
    for n, k in enumerate(liste):
        sirket = k["tur"] == "Kuruluş"
        bicimler = [(k["ad"], ad_kelimeleri(k["ad"], sirket))] + [(d, ad_kelimeleri(d, sirket)) for d in k["diger"]]
        bicimler = [(a, kel) for a, kel in bicimler if kel]
        hazir.append(bicimler)
        for _, kel in bicimler:
            for w in kel:
                indeks[w[:2]].add(n)
    sonuc = []
    for m in musteriler:
        sirket = m["tur"] == "Kuruluş"
        mk = ad_kelimeleri(m["ad"], sirket)
        adaylar = set().union(*(indeks.get(w[:2], set()) for w in mk)) if mk else set()
        for n in adaylar:
            k = liste[n]
            puan, hangi = 0.0, ""
            for ad, kel in hazir[n]:
                p = benzerlik(mk, kel)
                if p > puan:
                    puan, hangi = p, ad
            if puan < zayif:
                continue
            notlar, duzeltme = [], 0.0
            if (k["tur"] == "Kuruluş") != sirket:
                duzeltme -= 10
                notlar.append("kişi/kuruluş türü farklı")
            my = yil(m["dogum"])
            ly = {yil(d) for d in k["dogum"]} - {""}
            if my and ly:
                if my in ly:
                    tam = bool(iso(m["dogum"])) and iso(m["dogum"]) in {iso(d) for d in k["dogum"]}
                    duzeltme += 8 if tam else 5
                    notlar.append("doğum tarihi eşleşiyor" if tam else "doğum yılı eşleşiyor")
                elif min(abs(int(my) - int(y)) for y in ly) > 2:
                    duzeltme -= 15
                    notlar.append(f"doğum yılı farklı ({my} / {', '.join(sorted(ly))})")
            mu = katla(m["uyruk"])
            lu = {katla(u) for u in k["uyruk"]}
            if mu and lu:
                if any(mu == u or mu in u or u in mu for u in lu):
                    duzeltme += 3
                    notlar.append("uyruk eşleşiyor")
                else:
                    duzeltme -= 5
                    notlar.append("uyruk farklı")
            son = max(0.0, min(100.0, puan + duzeltme))
            seviye = ("Güçlü eşleşme" if son >= 95 else "Olası eşleşme" if son >= esik else "Zayıf benzerlik" if son >= zayif else None)
            if not seviye:
                continue
            sonuc.append({"musteri": m, "kayit": k, "ad_puan": round(puan, 1), "puan": round(son, 1), "eslesen_ad": hangi,
                          "seviye": seviye, "notlar": notlar})
    sonuc.sort(key=lambda x: (-x["puan"], x["musteri"]["no"]))
    return sonuc


def calistir(musteri_yolu: Path, liste_yollari: list[Path], cikti: Path, esik: float = 85.0, zayif: float = 75.0) -> dict:
    musteriler = musteriler_oku(musteri_yolu)
    liste = [k for y in liste_yollari for k in liste_oku(y)]
    if not liste:
        raise SystemExit("Yaptırım listesi boş.")
    eslesmeler = tara(musteriler, liste, esik, zayif)
    _rapor(musteriler, liste, liste_yollari, eslesmeler, esik, zayif, cikti)
    return {"musteriler": musteriler, "liste": liste, "eslesmeler": eslesmeler}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
SEVIYE_DOLGU = {"Güçlü eşleşme": "F8C9C6", "Olası eşleşme": "FFE8CC", "Zayıf benzerlik": "FFF4CE"}
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, satir=1):
    for c in ws[satir]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI


def _rapor(musteriler, liste, yollar, eslesmeler, esik, zayif, cikti):
    wb = Workbook()
    o = wb.active
    o.title = "Özet"
    o.append(["Yaptırım listesi taraması"])
    o["A1"].font = Font(bold=True, size=12)
    o.append(["Taranan kayıt", len(musteriler)])
    o.append(["Liste kaydı", f"{len(liste)} ({', '.join(y.name for y in yollar)})"])
    for s in ("Güçlü eşleşme", "Olası eşleşme", "Zayıf benzerlik"):
        o.append([s, sum(1 for e in eslesmeler if e["seviye"] == s)])
        o.cell(o.max_row, 1).fill = PatternFill("solid", fgColor=SEVIYE_DOLGU[s])
    o.append(["Eşleşmesi olan müşteri", len({e["musteri"]["no"] for e in eslesmeler})])
    o.append([])
    o.append(["Önemli", "Bu tarama bir ön elemedir; eşleşmeler kimlik bilgileri (doğum tarihi, uyruk, kimlik/vergi no, adres) ile "
                        "doğrulanmadan sonuç çıkarılmamalıdır. Gerçek eşleşmelerde ilgili mevzuat (ör. 6415 ve 7262 sayılı Kanunlar "
                        "kapsamındaki malvarlığı dondurma kararları) ve kurum prosedürü uygulanır. Listelerin güncelliği kullanıcının "
                        "sorumluluğundadır."])
    o.cell(o.max_row, 2).alignment = UST
    o.column_dimensions["A"].width = 26
    o.column_dimensions["B"].width = 110

    e = wb.create_sheet("Eşleşmeler")
    e.append(["Seviye", "Puan", "Ad Benzerliği", "Müşteri No", "Müşteri Ad / Unvan", "Rol", "Doğum", "Uyruk", "Listedeki Ad",
              "Eşleşen Biçim", "Liste", "Referans", "Listedeki Doğum", "Listedeki Uyruk", "Notlar", "Analist Kararı", "Açıklama"])
    _baslik(e)
    for x in eslesmeler:
        m, k = x["musteri"], x["kayit"]
        e.append([x["seviye"], x["puan"], x["ad_puan"], m["no"], m["ad"], m["rol"], m["dogum"], m["uyruk"], k["ad"], x["eslesen_ad"],
                  k["liste"], k["ref"], ", ".join(k["dogum"]), ", ".join(k["uyruk"]), "; ".join(x["notlar"]), None, None])
        e.cell(e.max_row, 1).fill = PatternFill("solid", fgColor=SEVIYE_DOLGU[x["seviye"]])
    for j, w in enumerate((15, 7, 9, 11, 28, 10, 11, 10, 30, 30, 18, 12, 14, 14, 36, 16, 30), 1):
        e.column_dimensions[get_column_letter(j)].width = w
    e.freeze_panes = "F2"
    e.auto_filter.ref = e.dimensions

    b = wb.create_sheet("Bilgi")
    for s in [["Normalleştirme", "Türkçe/aksanlı harfler, büyük/küçük harf, noktalama, unvanlar (Dr., Hacı, Al, Bin…) ve kuruluşlarda şirket ekleri "
                                 "(A.Ş., Ltd. Şti., LLC…) yok sayılır; yaygın ad transliterasyonları (Mohammed→Muhammed, Hussein→Hüseyin…) "
                                 "ve ph/f, w/v, q/k, kh/h, çift harf farkları eşitlenir"],
              ["Ad benzerliği", "kelimeler sıradan bağımsız eşlenir (en az %80 benzer kelimeler); puan = 2 × eşleşen benzerlik toplamı / "
                                "toplam kelime sayısı × 100; takma adların her biri ayrıca denenir"],
              ["Düzeltmeler", "doğum tarihi tam eşleşme +8, yıl eşleşmesi +5, yıl farkı 2'den büyük −15; uyruk eşleşmesi +3, farklı −5; "
                              "kişi/kuruluş türü farkı −10"],
              ["Seviyeler", f"≥ 95 güçlü eşleşme · ≥ {esik:g} olası eşleşme · ≥ {zayif:g} zayıf benzerlik"],
              ["Kaynaklar", "BM Güvenlik Konseyi Konsolide Listesi (resmî XML) doğrudan okunur; Resmî Gazete'de yayımlanan malvarlığı "
                            "dondurma kararları, OFAC, AB ve diğer listeler için Ad, Diğer Adlar, Doğum Tarihi, Uyruk, Liste, Referans "
                            "sütunlu bir CSV/Excel hazırlayın"]]:
        b.append(s)
    b.column_dimensions["A"].width = 16
    b.column_dimensions["B"].width = 130
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ornek = BURASI / "ornek_veri"
    ap = argparse.ArgumentParser(description="Müşteri/karşı taraf listesini yaptırım listeleriyle bulanık eşleştirir.")
    ap.add_argument("--musteriler", type=Path, default=ornek / "musteriler.csv",
                    help="Taranacak liste (.xlsx/.csv): Müşteri No, Ad Soyad / Unvan, Tip, Doğum Tarihi, Uyruk [, Rol]")
    ap.add_argument("--liste", type=Path, action="append",
                    help="Yaptırım listesi: BM konsolide XML veya CSV/Excel (Ad, Diğer Adlar, Doğum Tarihi, Uyruk, Liste, Referans); birden çok verilebilir")
    ap.add_argument("--esik", type=float, default=85.0, help="Olası eşleşme eşiği, 0–100 (varsayılan 85)")
    ap.add_argument("--zayif", type=float, default=75.0, help="Zayıf benzerlik eşiği (varsayılan 75)")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "yaptirim_taramasi.xlsx")
    a = ap.parse_args(argv)
    if not a.liste:
        a.liste = [ornek / "ornek_liste.xml", ornek / "ornek_liste.csv"]
    s = calistir(a.musteriler, a.liste, a.cikti, a.esik, a.zayif)
    from collections import Counter
    say = Counter(e["seviye"] for e in s["eslesmeler"])
    print(f"[OK] {len(s['musteriler'])} kayıt × {len(s['liste'])} liste kaydı · güçlü {say.get('Güçlü eşleşme', 0)} · "
          f"olası {say.get('Olası eşleşme', 0)} · zayıf {say.get('Zayıf benzerlik', 0)}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
