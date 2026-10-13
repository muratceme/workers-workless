import sys
import tempfile
import unittest
from datetime import date, datetime, time
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def d(makine, neden, sure, planli=False, vardiya="", satir=1):
    return {"tarih": date(2026, 10, 5), "makine": makine, "neden": neden, "sure": D(sure), "planli": planli, "vardiya": vardiya, "satir": satir}


def veri():
    return [d("M1", "Arıza", 100, vardiya="1", satir=2), d("M2", "arıza", 50, vardiya="2", satir=3), d("M1", "Ayar", 60, vardiya="1", satir=4), d("M2", "Malzeme bekleme", 30, vardiya="2", satir=5), d("M2", "Malzeme bekleme", 10, satir=6),
            d("M1", "Diğer", 10, vardiya="1", satir=7), d("M1", "Planlı bakım", 120, True, "1", 8)]


def k(l):
    return [(g["ad"], g["sinif"], g["adet"], g["sure"], g["pay"], g["kum"], g["ort"]) for g in l]


class AnalizTesti(unittest.TestCase):
    def test_pareto(self):
        r = main.analiz(veri())
        self.assertEqual(k(r["nedenler"]), [("Arıza", "A", 2, D(150), D("57.69"), D("57.69"), D("75.00")), ("Ayar", "A", 1, D(60), D("23.08"), D("80.77"), D("60.00")),
                                            ("Malzeme bekleme", "B", 2, D(40), D("15.38"), D("96.15"), D("20.00")), ("Diğer", "B", 1, D(10), D("3.85"), D("100.00"), D("10.00"))])
        self.assertEqual([(m["ad"], m["sinif"], m["adet"], m["sure"], m["pay"], m["kum"], m["en_neden"], m["en_pay"]) for m in r["makineler"]],
                         [("M1", "A", 3, D(170), D("65.38"), D("65.38"), "Arıza", D("58.82")), ("M2", "A", 3, D(90), D("34.62"), D("100.00"), "arıza", D("55.56"))])
        self.assertEqual([(v["ad"], v["adet"], v["sure"], v["pay"]) for v in r["vardiyalar"]], [("1", 3, D(170), D("68.00")), ("2", 2, D(80), D("32.00"))])
        self.assertEqual((r["ozet"], r["uyarilar"]), ({"adet": 6, "sure": D(260), "saat": D("4.33"), "ort": D("43.33"), "neden": 4, "a_neden": 2, "a_pay": D("80.77"), "planli_adet": 1, "planli_sure": D(120)}, []))

    def test_secenekler(self):
        r = main.analiz(veri(), esas="adet")
        self.assertEqual([(g["ad"], g["sinif"], g["pay"], g["kum"]) for g in r["nedenler"]], [("Arıza", "A", D("33.33"), D("33.33")), ("Malzeme bekleme", "A", D("33.33"), D("66.67")), ("Ayar", "A", D("16.67"), D("83.33")), ("Diğer", "B", D("16.67"), D("100.00"))])
        s = lambda **kw: [g["sinif"] for g in main.analiz(veri(), **kw)["nedenler"]]   # noqa: E731
        self.assertEqual((s(esik=D("57.69")), s(esik=D("57.70")), s(esik=D(100))), (["A", "B", "B", "B"], ["A", "A", "B", "B"], ["A"] * 4))
        r = main.analiz(veri(), planli_dahil=True, uzun=D(100))
        self.assertEqual(([g["ad"] for g in r["nedenler"]], r["ozet"]["sure"], r["ozet"]["adet"], [(u["tur"], u["kim"]) for u in r["uyarilar"]]), (["Arıza", "Planlı bakım", "Ayar", "Malzeme bekleme", "Diğer"], D(380), 7, [("Uzun duruş", "Satır 8 · M1")]))
        r = main.analiz([d("M", main.YOK, 50), d("M", "DİĞER", 30), d("M", "Arıza", 20)])
        self.assertEqual(sorted((u["onem"], u["tur"]) for u in r["uyarilar"]), [("Orta", "'Diğer' grubu büyük"), ("Orta", "Nedeni yazılmamış duruşlar")])
        for l in ([], [d("M", "P", 5, True)]):
            with self.assertRaises(ValueError):
                main.analiz(l)

    def test_dakika(self):
        self.assertEqual([main.dakika(x) for x in ("10:15", "9.05", "23:40:00", "05.10.2026 08:30", time(7, 45), datetime(2026, 10, 5, 22, 10), "24:00", "10:60", "x", "", None, "1015")], [615, 545, 1420, 510, 465, 1330, None, None, None, None, None, None])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "duruslar.csv", cikti)
            self.assertEqual(k(r["nedenler"]), [("Mekanik arıza", "A", 3, D(300), D("47.24"), D("47.24"), D("100.00")), ("Malzeme bekleme", "A", 3, D(100), D("15.75"), D("62.99"), D("33.33")),
                                                ("Kalıp değişimi", "A", 2, D(95), D("14.96"), D("77.95"), D("47.50")), ("Elektrik arızası", "A", 1, D(50), D("7.87"), D("85.83"), D("50.00")),
                                                ("Kalite onayı bekleme", "B", 1, D(30), D("4.72"), D("90.55"), D("30.00")), ("Takım kırılması", "B", 1, D(25), D("3.94"), D("94.49"), D("25.00")),
                                                ("Operatör yok", "B", 1, D(20), D("3.15"), D("97.64"), D("20.00")), ("Diğer", "B", 1, D(15), D("2.36"), D("100.00"), D("15.00"))])
            self.assertEqual([(m["ad"], m["adet"], m["sure"], m["en_neden"]) for m in r["makineler"]], [("Pres 1", 5, D(375), "Mekanik arıza"), ("CNC 1", 5, D(150), "Malzeme bekleme"), ("Pres 2", 3, D(110), "Elektrik arızası")])
            self.assertEqual([(v["ad"], v["sure"]) for v in r["vardiyalar"]], [("1", D(455)), ("2", D(130)), ("3", D(50))])
            o = r["ozet"]
            self.assertEqual((o["adet"], o["sure"], o["saat"], o["a_neden"], o["a_pay"], o["planli_adet"], o["planli_sure"]), (13, D(635), D("10.58"), 4, D("85.83"), 1, D(120)))
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]], [("Bilgi", "Uzun duruş", "Satır 5 · Pres 1")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Neden Pareto", "Makine Pareto", "Özet", "Duruşlar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "d.csv"
            y.write_text("Duruş Tarihi;Hat;Arıza Nedeni;Süre;Başlama;Bitiş Saati;Duruş Türü\n05.10.2026;H1;Arıza;12,5;;;PLANLI BAKIM\n;H1;;;22:00;01:30;\n;H2;X;;10:00;;\n;H2;X;0;;;\n;;X;5;;;\n;H3;Y;;08:00;08:00;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["tarih"], x["makine"], x["neden"], x["sure"], x["planli"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([(date(2026, 10, 5), "H1", "Arıza", D("12.5"), True), (None, "H1", main.YOK, D(210), False)], [("Duruş okunamadı", "Satır 4 · H2"), ("Duruş okunamadı", "Satır 5 · H2"), ("Duruş okunamadı", "Satır 7 · H3")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--duruslar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--esas", "adet", "--esik", "70", "--planli-dahil", "--uzun", "60"]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--esik", "0"], ["--esik", "101"], ["--uzun", "0"])], [2] * 3)
            self.assertEqual(main.main(["--duruslar", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Makine;Süre\nA;5\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            y.write_text("Makine;Duruş Nedeni;Süre;Tür\nA;B;5;Planlı\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
