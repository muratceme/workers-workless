"""
Duruş Nedeni Pareto Analizi — Gerisi Bende kod bloğu
Üretim › Üretim Mühendisi

Makine duruş kayıtlarını neden ve süreye göre Pareto analizine tabi tutar; duruşun çoğunu yaratan az sayıdaki nedeni öne çıkarır.
  Süre (dk)          = "Süre (dk)" sütunu; boşsa bitiş − başlangıç saati (bitiş başlangıçtan küçükse gece yarısını geçtiği kabul edilir)
  Pay %              = nedenin toplam süresi ÷ toplam duruş süresi × 100      (--esas adet ise adet üzerinden)
  Kümülatif %        = nedenler büyükten küçüğe sıralandığında payların birikimli toplamı
  A sınıfı neden     = kendisinden önceki kümülatif pay --esik (%80) değerinin altında olan neden
  Ortalama duruş     = süre ÷ adet
  Kapsam             = plansız duruşlar; --planli-dahil verilirse planlı duruşlar da girer
Duruş nedenini operatör kaydı belirler; kod kök neden analizi yapmaz.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
from datetime import time

SUTUN = {"tarih": ("tarih", "durus tarihi"), "makine": ("makine", "hat", "makine hat", "tezgah", "is merkezi"), "neden": ("durus nedeni", "neden", "durus kodu", "ariza nedeni"), "sure": ("sure dk", "sure", "durus suresi", "durus suresi dk"),
         "bas": ("baslangic", "baslangic saati", "baslama"), "bit": ("bitis", "bitis saati"), "tur": ("tur", "durus turu", "planli plansiz"), "vardiya": ("vardiya",)}
YOK = "(neden yazılmamış)"
DIGER = {"diger", "diger nedenler", "muhtelif", "bilinmiyor"}


def olcu(x) -> Decimal | None:
    s = metin(x).replace(" ", "")
    if not s:
        return None
    try:
        return Decimal(s.replace(".", "").replace(",", ".") if "," in s else s)
    except InvalidOperation:
        return None


def dakika(x) -> int | None:
    """Saat hücresini (10:15, 10.15, 10:15:00, datetime, time) gün içi dakikaya çevirir."""
    if isinstance(x, (datetime, time)):
        return x.hour * 60 + x.minute
    m = re.fullmatch(r"(\d{1,2})[:.](\d{2})(:\d{2})?", metin(x).split(" ")[-1])
    return int(m.group(1)) * 60 + int(m.group(2)) if m and int(m.group(1)) < 24 and int(m.group(2)) < 60 else None


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("makine", "neden")):
        makine = metin(r["makine"])
        if not makine:
            continue
        kim = f"Satır {r['_satir']} · {makine}"
        sure, bas, bit = para(r["sure"]), dakika(r["bas"]), dakika(r["bit"])
        if sure is None and bas is not None and bit is not None:
            sure = Decimal(bit - bas + (1440 if bit < bas else 0))
        if sure is None or sure <= 0:
            uy.append(uyari("Yüksek", "Duruş okunamadı", kim, f"Süre '{metin(r['sure'])}' okunamadı ve başlangıç–bitiş saatlerinden de hesaplanamadı. Satır alınmadı."))
            continue
        l.append({"tarih": tarih(r["tarih"]), "makine": makine, "neden": metin(r["neden"]) or YOK, "sure": sure, "planli": katla(r["tur"]).startswith("planli"), "vardiya": metin(r["vardiya"]), "satir": r["_satir"]})
    return l, uy


def _oran(pay, payda) -> Decimal | None:
    return yuvarla(Decimal(pay) * 100 / Decimal(payda)) if payda else None


def pareto(kayit: list[dict], alan: str, esas: str = "sure", esik: Decimal = Decimal(80)) -> list[dict]:
    grup = {}
    for d in kayit:
        g = grup.setdefault(katla(d[alan]), {"ad": d[alan], "adet": 0, "sure": Decimal(0)})
        g["adet"], g["sure"] = g["adet"] + 1, g["sure"] + d["sure"]
    s = sorted(grup.values(), key=lambda g: (-g[esas], -g["sure"], katla(g["ad"])))
    toplam, kum = sum(Decimal(g[esas]) for g in s), Decimal(0)
    for g in s:
        g["ort"], g["pay"] = yuvarla(g["sure"] / g["adet"]), _oran(g[esas], toplam)
        g["sinif"] = "A" if _oran(kum, toplam) < esik else "B"
        kum += Decimal(g[esas])
        g["kum"] = _oran(kum, toplam)
    return s


def analiz(duruslar: list[dict], esas: str = "sure", esik: Decimal = Decimal(80), planli_dahil: bool = False, uzun: Decimal = Decimal(120)) -> dict:
    kapsam = [d for d in duruslar if planli_dahil or not d["planli"]]
    if not kapsam:
        raise ValueError("Analiz edilecek duruş kaydı yok.")
    uy = []
    nedenler = pareto(kapsam, "neden", esas, esik)
    makineler = pareto(kapsam, "makine", esas, esik)
    for m in makineler:
        en = pareto([d for d in kapsam if katla(d["makine"]) == katla(m["ad"])], "neden", esas, esik)[0]
        m["en_neden"], m["en_pay"] = en["ad"], en["pay"]
    for n in nedenler:
        if n["ad"] == YOK:
            uy.append(uyari("Orta", "Nedeni yazılmamış duruşlar", YOK, f"{n['adet']} duruş, {n['sure']} dk (%{n['pay']}). Neden yazılmadan Pareto analizi eksik kalır."))
        elif katla(n["ad"]) in DIGER and n["pay"] > 10:
            uy.append(uyari("Orta", "'Diğer' grubu büyük", n["ad"], f"Toplamın %{n['pay']} kadarı ayrıştırılmamış nedenlerde ({n['adet']} duruş, {n['sure']} dk). Bu grubu alt nedenlere bölün."))
    for d in kapsam:
        if d["sure"] > uzun:
            uy.append(uyari("Bilgi", "Uzun duruş", f"Satır {d['satir']} · {d['makine']}", f"{d['neden']}: {d['sure']} dk" + (f" ({d['tarih']:%d.%m.%Y})" if d["tarih"] else "") + f"; tek duruşta {uzun} dakikayı aşıyor."))
    a = [n for n in nedenler if n["sinif"] == "A"]
    sure = sum((d["sure"] for d in kapsam), Decimal(0))
    return {"duruslar": kapsam, "nedenler": nedenler, "makineler": makineler, "vardiyalar": sorted(pareto([d for d in kapsam if d["vardiya"]], "vardiya", esas, esik), key=lambda v: katla(v["ad"])),
            "ozet": {"adet": len(kapsam), "sure": sure, "saat": yuvarla(sure / 60), "ort": yuvarla(sure / len(kapsam)), "neden": len(nedenler), "a_neden": len(a), "a_pay": a[-1]["kum"], "planli_adet": sum(1 for d in duruslar if d["planli"]),
                     "planli_sure": sum((d["sure"] for d in duruslar if d["planli"]), Decimal(0))}, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, esas: str, esik: Decimal, planli_dahil: bool) -> None:
    wb = Workbook()
    olcut = "süre" if esas == "sure" else "adet"
    for i, (baslik, alan, ad) in enumerate((("Neden Pareto", "nedenler", "Duruş Nedeni"), ("Makine Pareto", "makineler", "Makine / Hat"))):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = baslik
        _baslik(ws, ["Sıra", ad, "Sınıf", "Duruş Adedi", "Toplam Süre (dk)", f"Pay % ({olcut})", "Kümülatif %", "Ortalama Duruş (dk)"] + (["En Büyük Nedeni", "Nedenin Makinedeki Payı %"] if i else []), (6, 34, 7, 10, 13, 11, 12, 13, 30, 12))
        for n, g in enumerate(r[alan], 1):
            _satir(ws, [n, g["ad"], g["sinif"], g["adet"], g["sure"], g["pay"], g["kum"], g["ort"]] + ([g["en_neden"], g["en_pay"]] if i else []), para_s=(6, 7, 8, 10), dolgu={3: KIRMIZI if g["sinif"] == "A" else None})
        ws.append([])
        ws.append([f"A sınıfı: toplamın ilk %{esik} kadarını oluşturan kalemler. Sıralama ölçütü: {olcut}."])
    ws = wb.create_sheet("Özet")
    o = r["ozet"]
    _baslik(ws, ["Gösterge", "Değer"], (48, 16))
    for etiket, deger, para_mi in (("Duruş adedi", o["adet"], False), ("Toplam duruş süresi (dk)", o["sure"], True), ("Toplam duruş süresi (saat)", o["saat"], True), ("Ortalama duruş süresi (dk)", o["ort"], True),
                                   ("Farklı neden sayısı", o["neden"], False), ("A sınıfı neden sayısı", o["a_neden"], False), ("A sınıfı nedenlerin toplamdaki payı %", o["a_pay"], True),
                                   ("Planlı duruş adedi" + ("" if planli_dahil else " (analize girmedi)"), o["planli_adet"], False), ("Planlı duruş süresi (dk)" + ("" if planli_dahil else " (analize girmedi)"), o["planli_sure"], True)):
        _satir(ws, [etiket, deger], para_s=(2,) if para_mi else ())
    if r["vardiyalar"]:
        ws.append([])
        _baslik(ws, ["Vardiya", "Duruş Adedi", "Toplam Süre (dk)", "Pay %", "Ortalama Duruş (dk)"], dondur="A2")
        for v in r["vardiyalar"]:
            _satir(ws, [v["ad"], v["adet"], v["sure"], v["pay"], v["ort"]], para_s=(4, 5))
        for h in "CDE":
            ws.column_dimensions[h].width = 16
    ws = wb.create_sheet("Duruşlar")
    _baslik(ws, ["Tarih", "Vardiya", "Makine / Hat", "Duruş Nedeni", "Süre (dk)", "Tür"], (12, 10, 22, 34, 11, 10))
    for d in sorted(r["duruslar"], key=lambda d: (d["tarih"] or date.min, d["satir"])):
        _satir(ws, [d["tarih"], d["vardiya"], d["makine"], d["neden"], d["sure"], "Planlı" if d["planli"] else "Plansız"], tarih_s=(1,))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, esas: str = "sure", esik: Decimal = Decimal(80), planli_dahil: bool = False, uzun: Decimal = Decimal(120)) -> dict:
    duruslar, uy = oku(yol)
    r = analiz(duruslar, esas, esik, planli_dahil, uzun)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, esas, esik, planli_dahil)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Makine duruş kayıtlarını neden ve süreye göre Pareto analizine tabi tutar.")
    p.add_argument("--duruslar", type=Path, help="Tarih, Vardiya, Makine, Duruş Nedeni, Süre (dk) ya da Başlangıç ve Bitiş, Tür (Planlı/Plansız)")
    p.add_argument("--esas", choices=("sure", "adet"), default="sure", help="Sıralama ölçütü: sure (varsayılan) ya da adet")
    p.add_argument("--esik", default="80", help="A sınıfı için kümülatif pay sınırı, yüzde (varsayılan 80)")
    p.add_argument("--planli-dahil", action="store_true", help="Planlı duruşları da analize kat")
    p.add_argument("--uzun", default="120", help="Tek duruşta bu dakikayı aşanlar ayrıca bildirilir (varsayılan 120)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "durus_pareto_analizi.xlsx")
    a = p.parse_args(argv)
    yol = a.duruslar or ORNEK / "duruslar.csv"
    esik, uzun = olcu(a.esik), olcu(a.uzun)
    if esik is None or not 0 < esik <= 100 or uzun is None or uzun <= 0:
        print("[X] --esik 0 ile 100 arasında; --uzun sıfırdan büyük olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, a.esas, esik, a.planli_dahil, uzun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {o['adet']} duruş · {o['sure']} dk ({o['saat']} saat) · {o['neden']} neden · A sınıfı {o['a_neden']} neden toplamın %{o['a_pay']} kadarı")
    for n in r["nedenler"]:
        if n["sinif"] == "A":
            print(f"     {n['ad']}: {n['sure']} dk · {n['adet']} duruş · %{n['pay']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
