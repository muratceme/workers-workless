import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
KUR = {"USD": D(40), "EUR": D("45.5")}


def k(tur, pb, tutar):
    return {"kalem": "K", "tur": tur, "pb": pb, "tutar": D(str(tutar))}


class HesapTesti(unittest.TestCase):
    def test_pozisyon(self):
        l = [k("varlik", "USD", 90), k("varlik", "USD", 10), k("yukumluluk", "USD", 130), k("alim", "USD", 50), k("satim", "USD", 5), k("varlik", "EUR", 100), k("yukumluluk", "EUR", 60), k("satim", "EUR", 40), k("yukumluluk", "CHF", 7),
             k("varlik", "TRY", 999), k("alim", "TL", 1)]
        r = main.hesapla(l, KUR, sok=D(20))
        self.assertEqual([(p["pb"], p["varlik"], p["yukumluluk"], p["ici"], p["alim"], p["satim"], p["disi"], p["net"], p["yon"], p["kur"], p["ici_tl"], p["disi_tl"], p["net_tl"], p["sok_etki"]) for p in r["pozisyonlar"]],
                         [("CHF", D(0), D(7), D(-7), D(0), D(0), D(0), D(-7), "Kısa", None, None, None, None, None), ("EUR", D(100), D(60), D(40), D(0), D(40), D(-40), D(0), "Kapalı", D("45.5"), D(1820), D(-1820), D(0), D(0)),
                          ("USD", D(100), D(130), D(-30), D(50), D(5), D(45), D(15), "Uzun", D(40), D(-1200), D(1800), D(600), D(120))])
        self.assertEqual(r["toplam"], {"uzun": D(600), "kisa": D(0), "ici": D(620), "disi": D(-20), "net": D(600), "yon": "Uzun", "oran": None, "sok_etki": D(120)})
        self.assertEqual((sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]), len(r["kalemler"])), ([("Bilgi", "TL kalem dikkate alınmadı", "2 satır"), ("Yüksek", "Kur yok", "CHF")], 9))
        for bos in ([], [k("varlik", "TRY", 5)]):
            with self.assertRaises(ValueError):
                main.hesapla(bos, KUR)

    def test_ozkaynak_ve_limit(self):
        kisa = [k("yukumluluk", "USD", 100), k("varlik", "EUR", 20)]   # −4000 + 910 = −3090
        t = lambda l, **kw: (lambda r: (r["toplam"]["net"], r["toplam"]["yon"], r["toplam"]["oran"], r["toplam"]["sok_etki"], [(u["onem"], u["tur"]) for u in r["uyarilar"]]))(main.hesapla(l, KUR, **kw))   # noqa: E731
        self.assertEqual(t(kisa, ozkaynak=D(30900)), (D(-3090), "Kısa", D("-10.00"), D(-309), []))
        self.assertEqual(t(kisa, ozkaynak=D(30900), limit=D(10))[4], [("Orta", "Pozisyon limitine yaklaştı")])
        self.assertEqual(t(kisa, ozkaynak=D(30900), limit=D("9.99"))[4], [("Yüksek", "Pozisyon limiti aşıldı")])
        self.assertEqual((t(kisa, ozkaynak=D(30900), limit=D("12.5"))[4], t(kisa, ozkaynak=D(30900), limit=D("12.51"))[4]), ([("Orta", "Pozisyon limitine yaklaştı")], []))
        self.assertEqual(t([k("varlik", "USD", 100)], ozkaynak=D(20000), limit=D(15)), (D(4000), "Uzun", D("20.00"), D(400), [("Yüksek", "Pozisyon limiti aşıldı")]))
        r = main.hesapla([k("varlik", "USD", 10), k("yukumluluk", "EUR", 5)], KUR)["toplam"]
        self.assertEqual((r["uzun"], r["kisa"], r["net"]), (D(400), D("227.5"), D("172.5")))
        self.assertEqual(main.hesapla([k("varlik", "USD", 1), k("yukumluluk", "USD", 1)], KUR)["toplam"]["yon"], "Kapalı")


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "kalemler.csv", ORNEK / "kurlar.csv", cikti, D(2000000000))
            m = 1000000
            self.assertEqual([(p["pb"], p["varlik"], p["yukumluluk"], p["ici"], p["disi"], p["net"], p["yon"], p["ici_tl"], p["disi_tl"], p["net_tl"], p["sok_etki"]) for p in r["pozisyonlar"]],
                             [("EUR", D(45 * m), D(38 * m), D(7 * m), D(-9 * m), D(-2 * m), "Kısa", D(322 * m), D(-414 * m), D(-92 * m), D(-9200000)), ("GBP", D(m), D(1500000), D(-500000), D(0), D(-500000), "Kısa", D(-27 * m), D(0), D(-27 * m), D(-2700000)),
                              ("USD", D(90 * m), D(100 * m), D(-10 * m), D(12 * m), D(2 * m), "Uzun", D(-425 * m), D(510 * m), D(85 * m), D(8500000))])
            self.assertEqual(r["toplam"], {"uzun": D(85 * m), "kisa": D(119 * m), "ici": D(-130 * m), "disi": D(96 * m), "net": D(-34 * m), "yon": "Kısa", "oran": D("-1.70"), "sok_etki": D(-3400000)})
            self.assertEqual((r["uyarilar"], load_workbook(cikti).sheetnames), ([], ["Pozisyon", "Kalemler", "Uyarılar"]))

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            ky, ry = Path(tmp) / "k.csv", Path(tmp) / "r.csv"
            ky.write_text("Hesap;Sınıf;Döviz;Bakiye\nKrediler;AKTİF;usd;1.250.000,50\n;Türev Satım;Eur;5\nX;Özkaynak;USD;5\nY;Varlık;;5\nZ;Varlık;USD;-1\n;;;\n", encoding="utf-8")
            ry.write_text("Döviz Cinsi;Değerleme Kuru\nusd;42,5\nEUR;46.25\n;3\n", encoding="utf-8")
            l, uy = main.oku_kalemler(ky)
            self.assertEqual((l, [(u["tur"], u["kim"]) for u in uy]),
                             ([{"kalem": "Krediler", "tur": "varlik", "pb": "USD", "tutar": D("1250000.50")}, {"kalem": "", "tur": "satim", "pb": "EUR", "tutar": D(5)}], [("Kalem okunamadı", n) for n in "XYZ"]))
            self.assertEqual(main.oku_kurlar(ry), {"USD": D("42.5"), "EUR": D("46.25")})
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--kalemler", str(ky), "--kurlar", str(ry), "--cikti", c]
            self.assertEqual(main.main(arg + ["--ozkaynak", "2.000.000.000", "--limit", "20", "--sok", "15"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--ozkaynak", "0"], ["--limit", "20"], ["--ozkaynak", "5", "--limit", "0"], ["--sok", "0"], ["--sok", "x"])] + [main.main(["--kalemler", str(ky)])], [2] * 6)
            self.assertEqual(main.main(["--kalemler", str(ky), "--kurlar", str(Path(tmp) / "yok.csv")]), 1)
            ry.write_text("Para Birimi;Kur\nUSD;0\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            ry.write_text("Para Birimi;Kur\nUSD;40\n", encoding="utf-8")
            for icerik in ("Kalem;Tutar\nA;1\n", "Tür;Para Birimi;Tutar\n;;\n", "Tür;Para Birimi;Tutar\nVarlık;TL;5\n"):
                ky.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
