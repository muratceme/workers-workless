"""
Döviz Pozisyonu Raporu — Workers / Workless kod bloğu
Hazine ve Aktif-Pasif Yönetimi › Hazine Uzmanı

Bilanço içi ve dışı döviz varlık ve yükümlülüklerinden döviz cinsi bazında net pozisyonu hesaplar.
  Bilanço içi net       = döviz varlıkları − döviz yükümlülükleri
  Bilanço dışı net      = döviz alım taahhütleri − döviz satım taahhütleri   (vadeli, swap, opsiyon vb.)
  Net genel pozisyon    = bilanço içi net + bilanço dışı net   (artı: uzun, eksi: kısa)
  TL karşılığı          = pozisyon × kur
  Uzun / kısa toplam    = TL karşılığı artı olan ve eksi olan pozisyonların ayrı ayrı toplamı
  Pozisyon ÷ özkaynak % = net genel pozisyon toplamı (TL) ÷ özkaynak × 100   (--ozkaynak verilirse)
  Kur şoku etkisi       = pozisyonun TL karşılığı × şok yüzdesi ÷ 100   (TL'nin o para birimine karşı değer kaybı)
Tutarlar orijinal para biriminde, kurlar TL karşılığı olarak girilir. Yasal oran hesabındaki özel kurallar (endeksli kalemler, opsiyon deltası, istisnalar) uygulanmaz.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
K_SUTUN = {"kalem": ("kalem", "hesap", "aciklama"), "tur": ("tur", "kalem turu", "sinif"), "pb": ("para birimi", "doviz", "doviz cinsi", "pb"), "tutar": ("tutar", "bakiye")}
KUR_SUTUN = {"pb": ("para birimi", "doviz", "doviz cinsi", "pb"), "kur": ("kur", "tl karsiligi", "degerleme kuru")}
TUR = {"bilanco ici varlik": "varlik", "varlik": "varlik", "aktif": "varlik", "bilanco ici yukumluluk": "yukumluluk", "yukumluluk": "yukumluluk", "pasif": "yukumluluk", "bilanco disi alim": "alim", "turev alim": "alim", "alim": "alim",
       "vadeli alim": "alim", "bilanco disi satim": "satim", "turev satim": "satim", "satim": "satim", "vadeli satim": "satim"}
TL = {"TRY", "TL", "TURK LIRASI"}
Z, YUZ = Decimal(0), Decimal(100)


def olcu(x) -> Decimal | None:
    """Kur okur: virgül ya da tek nokta ondalık ayırıcıdır (42,50 = 42.50)."""
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x or ""))
    try:
        return Decimal(s.replace(".", "").replace(",", ".") if "," in s else s) if s else None
    except InvalidOperation:
        return None


def oku_kalemler(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, K_SUTUN, ("tur", "pb", "tutar")):
        if not metin(r["tur"]) and not metin(r["pb"]):
            continue
        kim = metin(r["kalem"]) or f"Satır {r['_satir']}"
        tur, pb, tutar = TUR.get(katla(r["tur"])), katla(r["pb"]).upper(), para(r["tutar"])
        if tur is None or not pb or tutar is None or tutar < 0:
            uy.append(uyari("Yüksek", "Kalem okunamadı", kim, f"Satır {r['_satir']}: tür '{metin(r['tur'])}' (Bilanço İçi Varlık / Bilanço İçi Yükümlülük / Bilanço Dışı Alım / Bilanço Dışı Satım olmalı), para birimi '{metin(r['pb'])}', "
                            f"tutar '{metin(r['tutar'])}'. Satır alınmadı."))
            continue
        l.append({"kalem": metin(r["kalem"]), "tur": tur, "pb": pb, "tutar": tutar})
    return l, uy


def oku_kurlar(yol: Path) -> dict[str, Decimal]:
    kurlar = {}
    for r in kayitlar(yol, KUR_SUTUN, ("pb", "kur")):
        if not metin(r["pb"]):
            continue
        kur = olcu(r["kur"])
        if kur is None or kur <= 0:
            raise ValueError(f"{yol.name} satır {r['_satir']}: kur '{metin(r['kur'])}' pozitif bir sayı olmalı.")
        kurlar[katla(r["pb"]).upper()] = kur
    return kurlar


def hesapla(kalemler: list[dict], kurlar: dict[str, Decimal], ozkaynak: Decimal | None = None, limit: Decimal | None = None, sok: Decimal = Decimal(10)) -> dict:
    uy, pbler = [], {}
    tl_kalem = [k for k in kalemler if k["pb"] in TL]
    if tl_kalem:
        uy.append(uyari("Bilgi", "TL kalem dikkate alınmadı", f"{len(tl_kalem)} satır", "Para birimi TL olan kalemler döviz pozisyonuna girmez."))
    for k in kalemler:
        if k["pb"] not in TL:
            p = pbler.setdefault(k["pb"], {"pb": k["pb"], "varlik": Z, "yukumluluk": Z, "alim": Z, "satim": Z})
            p[k["tur"]] += k["tutar"]
    if not pbler:
        raise ValueError("Döviz kalemi yok.")
    liste = []
    for pb, p in sorted(pbler.items()):
        p["ici"], p["disi"] = p["varlik"] - p["yukumluluk"], p["alim"] - p["satim"]
        p["net"] = p["ici"] + p["disi"]
        p["kur"] = kurlar.get(pb)
        p["yon"] = "Uzun" if p["net"] > 0 else "Kısa" if p["net"] < 0 else "Kapalı"
        if p["kur"] is None:
            uy.append(uyari("Yüksek", "Kur yok", pb, f"Net pozisyon {p['net']:,.2f} {pb} için kur verilmedi; TL karşılığı hesaplanamadı ve toplamlara katılmadı."))
        for alan in ("ici", "disi", "net"):
            p[alan + "_tl"] = yuvarla(p[alan] * p["kur"]) if p["kur"] else None
        p["sok_etki"] = yuvarla(p["net_tl"] * sok / YUZ) if p["kur"] else None
        liste.append(p)
    kurlu = [p for p in liste if p["kur"]]
    t = {"uzun": sum((p["net_tl"] for p in kurlu if p["net_tl"] > 0), Z), "kisa": sum((-p["net_tl"] for p in kurlu if p["net_tl"] < 0), Z), "ici": sum((p["ici_tl"] for p in kurlu), Z), "disi": sum((p["disi_tl"] for p in kurlu), Z)}
    t["net"] = t["uzun"] - t["kisa"]
    t["yon"] = "Uzun" if t["net"] > 0 else "Kısa" if t["net"] < 0 else "Kapalı"
    t["oran"] = yuvarla(t["net"] * YUZ / ozkaynak) if ozkaynak else None
    t["sok_etki"] = yuvarla(t["net"] * sok / YUZ)
    if limit is not None and ozkaynak and abs(t["net"] * YUZ / ozkaynak) > limit:
        uy.append(uyari("Yüksek", "Pozisyon limiti aşıldı", "Net genel pozisyon", f"Net genel pozisyon {t['net']:,.2f} TL ({t['yon'].lower()}), özkaynağın %{abs(t['oran'])}'i; limit %{limit:f}."))
    elif limit is not None and ozkaynak and abs(t["net"] * YUZ / ozkaynak) >= limit * Decimal("0.8"):
        uy.append(uyari("Orta", "Pozisyon limitine yaklaştı", "Net genel pozisyon", f"Net genel pozisyon {t['net']:,.2f} TL ({t['yon'].lower()}), özkaynağın %{abs(t['oran'])}'i; limit %{limit:f}."))
    return {"pozisyonlar": liste, "toplam": t, "kalemler": [k for k in kalemler if k["pb"] not in TL], "ozkaynak": ozkaynak, "limit": limit, "sok": sok, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Pozisyon"
    sok = r["sok"]
    _baslik(ws, ["Para Birimi", "Varlıklar", "Yükümlülükler", "Bilanço İçi Net", "Bilanço Dışı Alım", "Bilanço Dışı Satım", "Bilanço Dışı Net", "Net Genel Pozisyon", "Yön", "Kur", "Bilanço İçi Net (TL)", "Bilanço Dışı Net (TL)",
                 "Net Genel Pozisyon (TL)", f"Kur %{sok:f} Artarsa Etki (TL)"], (10, 17, 17, 17, 17, 17, 17, 17, 9, 10, 19, 19, 19, 19))
    for p in r["pozisyonlar"]:
        _satir(ws, [p["pb"], p["varlik"], p["yukumluluk"], p["ici"], p["alim"], p["satim"], p["disi"], p["net"], p["yon"], p["kur"], p["ici_tl"], p["disi_tl"], p["net_tl"], p["sok_etki"]], para_s=(2, 3, 4, 5, 6, 7, 8, 11, 12, 13, 14),
               dolgu={9: {"Uzun": YESIL, "Kısa": SARI}.get(p["yon"])})
        ws.cell(ws.max_row, 10).number_format = "#,##0.0000"
    t = r["toplam"]
    n = _satir(ws, ["TOPLAM (TL)", None, None, None, None, None, None, None, t["yon"], None, t["ici"], t["disi"], t["net"], t["sok_etki"]], para_s=(11, 12, 13, 14))
    for h in ws[n]:
        h.font = Font(bold=True)
    ws.append([])
    for etiket, deger in (("Uzun pozisyonlar toplamı (TL)", t["uzun"]), ("Kısa pozisyonlar toplamı (TL)", t["kisa"]), ("Net genel pozisyon (TL)", t["net"]), ("Özkaynak (TL)", r["ozkaynak"]), ("Net genel pozisyon ÷ özkaynak %", t["oran"]),
                          ("Limit %", r["limit"])):
        if deger is not None:
            _satir(ws, [etiket, None, None, None, None, None, None, None, None, None, None, None, deger], para_s=(13,))
    ws.append([])
    ws.append(["Artı pozisyon uzun (döviz fazlası), eksi pozisyon kısa (döviz açığı) demektir. Kur şoku etkisi, TL o para birimine karşı değer kaybederse oluşacak kazanç (artı) ya da zarardır (eksi)."])
    ws = wb.create_sheet("Kalemler")
    _baslik(ws, ["Kalem", "Tür", "Para Birimi", "Tutar"], (36, 26, 10, 18))
    adlar = {"varlik": "Bilanço İçi Varlık", "yukumluluk": "Bilanço İçi Yükümlülük", "alim": "Bilanço Dışı Alım", "satim": "Bilanço Dışı Satım"}
    for k in sorted(r["kalemler"], key=lambda k: (k["pb"], list(adlar).index(k["tur"]))):
        _satir(ws, [k["kalem"], adlar[k["tur"]], k["pb"], k["tutar"]], para_s=(4,))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(k_yol: Path, kur_yol: Path, cikti: Path, ozkaynak: Decimal | None = None, limit: Decimal | None = None, sok: Decimal = Decimal(10)) -> dict:
    kalemler, uy = oku_kalemler(k_yol)
    r = hesapla(kalemler, oku_kurlar(kur_yol), ozkaynak, limit, sok)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Bilanço içi ve dışı döviz kalemlerinden döviz cinsi bazında net pozisyonu hesaplar.")
    p.add_argument("--kalemler", type=Path, help="Kalem, Tür (Bilanço İçi Varlık / Bilanço İçi Yükümlülük / Bilanço Dışı Alım / Bilanço Dışı Satım), Para Birimi, Tutar")
    p.add_argument("--kurlar", type=Path, help="Para Birimi, Kur (TL karşılığı)")
    p.add_argument("--ozkaynak", help="Özkaynak (TL); verilirse pozisyonun özkaynağa oranı hesaplanır")
    p.add_argument("--limit", help="Net genel pozisyon / özkaynak için limit, yüzde (mutlak değer; verilmezse karşılaştırma yapılmaz)")
    p.add_argument("--sok", default="10", help="Kur şoku senaryosu, yüzde (varsayılan 10)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "doviz_pozisyonu.xlsx")
    a = p.parse_args(argv)
    if (a.kalemler is None) != (a.kurlar is None):
        print("[X] --kalemler ve --kurlar birlikte verilmelidir")
        return 2
    ornek = a.kalemler is None
    k_yol, kur_yol = (ORNEK / "kalemler.csv", ORNEK / "kurlar.csv") if ornek else (a.kalemler, a.kurlar)
    ozkaynak, limit, sok = para(a.ozkaynak or ("2000000000" if ornek else None)), olcu(a.limit), olcu(a.sok)
    if ((a.ozkaynak or ornek) and (ozkaynak is None or ozkaynak <= 0)) or (a.limit and (limit is None or limit <= 0)) or (a.limit and not ozkaynak) or sok is None or sok <= 0:
        print("[X] --ozkaynak pozitif; --limit pozitif ve --ozkaynak ile birlikte; --sok pozitif olmalı")
        return 2
    for y in (k_yol, kur_yol):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(k_yol, kur_yol, a.cikti, ozkaynak, limit, sok)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {len(r['pozisyonlar'])} para birimi · uzun {t['uzun']:,.2f} TL · kısa {t['kisa']:,.2f} TL · net genel pozisyon {t['net']:,.2f} TL ({t['yon'].lower()})" + (f" · özkaynağın %{t['oran']}'i" if t["oran"] is not None else ""))
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
