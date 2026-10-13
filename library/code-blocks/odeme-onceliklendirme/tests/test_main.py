import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BAS = date(2026, 10, 12)


def k(cari, tutar, vade=date(2026, 10, 14), onem=2, faiz=0, zorunlu=False):
    return {"cari": cari, "belge": "", "vade": vade, "tutar": D(tutar), "tur": "", "onem": onem, "faiz": D(faiz), "zorunlu": zorunlu, "kim": cari}


class SiralamaTesti(unittest.TestCase):
    def test_sira_kurali(self):
        l = [k("tutar-büyük", 300), k("tutar-küçük", 100), k("vade-erken", 300, date(2026, 10, 13)), k("maliyetli", 3000, faiz=3), k("gecikmiş", 500, date(2026, 10, 1)), k("düşük-gecikmiş", 10, date(2026, 9, 1), onem=3),
             k("yüksek", 900, date(2026, 10, 18), onem=1), k("kritik", 900, date(2026, 10, 18), onem=0), k("zorunlu", 900, date(2026, 10, 18), zorunlu=True), k("ileri", 5, date(2026, 10, 19), onem=0, zorunlu=True)]
        r = main.onceliklendir(l, D(100000), BAS)
        self.assertEqual([x["cari"] for x in r["sirali"]], ["zorunlu", "kritik", "yüksek", "gecikmiş", "maliyetli", "vade-erken", "tutar-küçük", "tutar-büyük", "düşük-gecikmiş"])
        self.assertEqual([(x["cari"], x["karar"], x["sira"]) for x in r["ileri"]], [("ileri", "Vadesi gelmedi", None)])
        self.assertEqual([(x["grup"], x["gecikme"], x["gunluk"]) for x in l[3:6]], [("Dönem içi", 0, D("3.00")), ("Vadesi geçmiş", 11, D("0.00")), ("Vadesi geçmiş", 41, D("0.00"))])
        self.assertEqual((r["son"], r["ihtiyac"], r["odenecek"], r["ertelenen"], r["kalan"], r["ileri_toplam"], r["uyarilar"]), (date(2026, 10, 18), D(6910), D(6910), D(0), D(93090), D(5), []))

    def test_dagitim(self):
        l = [k("A", 400, zorunlu=True), k("B", 500, onem=0), k("C", 300, date(2026, 10, 1), onem=1, faiz=3), k("D", 100, date(2026, 10, 1)), k("E", 50, onem=3)]
        r = main.onceliklendir(l, D(700), BAS)
        self.assertEqual([(x["cari"], x["karar"], x["kalan_nakit"]) for x in r["sirali"]], [("A", "Öde", D(300)), ("B", "Ertele", D(300)), ("C", "Öde", D(0)), ("D", "Ertele", D(0)), ("E", "Ertele", D(0))])
        self.assertEqual((r["ihtiyac"], r["odenecek"], r["ertelenen"], r["kalan"], r["ertelenen_gunluk"]), (D(1350), D(700), D(650), D(0), D(0)))
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]), [("Orta", "Ertelenen ödeme", "D"), ("Orta", "Ertelenen ödeme", "E"), ("Yüksek", "Kritik ödeme erteleniyor", "B"), ("Yüksek", "Nakit yetersiz", "Dönem")])
        self.assertEqual(r["sirali"][2]["gerekce"], "Yüksek · 11 gün gecikmiş · günlük gecikme maliyeti 0.30")
        r = main.onceliklendir([k("A", 400, zorunlu=True), k("C", 300, date(2026, 10, 1), faiz=3), k("E", 50)], D(700), BAS, 7, D(350))
        self.assertEqual(([(x["cari"], x["karar"]) for x in r["sirali"]], r["kullanilabilir"], r["kalan"], r["ertelenen_gunluk"], sorted(u["tur"] for u in r["uyarilar"])),
                         ([("A", "Ertele"), ("C", "Öde"), ("E", "Öde")], D(350), D(0), D(0), ["Nakit yetersiz", "Zorunlu ödeme karşılanamıyor"]))
        r = main.onceliklendir([k("C", 300, date(2026, 10, 1), faiz=3)], D(100), BAS, 7, D(200))
        self.assertEqual((r["sirali"][0]["karar"], r["kalan"], r["ertelenen_gunluk"]), ("Ertele", D(-100), D("0.30")))
        self.assertEqual(main.onceliklendir([k("X", 5, date(2026, 10, 13))], D(10), BAS, 1)["ileri_toplam"], D(5))
        with self.assertRaises(ValueError):
            main.onceliklendir([], D(1), BAS)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "o.xlsx"
            r = main.calistir(ORNEK / "odemeler.csv", cikti, D(900000), BAS)
            self.assertEqual([(x["cari"], x["karar"], x["kalan_nakit"]) for x in r["sirali"]],
                             [("A Bankası", "Öde", D(720000)), ("Personel", "Öde", D(300000)), ("Enerji Dağıtım", "Öde", D(228000)), ("Demir Çelik A.Ş.", "Öde", D(78000)), ("Ofis Kira", "Öde", D(33000)),
                              ("Lojistik A.Ş.", "Ertele", D(33000)), ("Ambalaj Ltd.", "Ertele", D(33000)), ("Kırtasiye Ltd.", "Öde", D(25000))])
            self.assertEqual(([x["cari"] for x in r["ileri"]], r["ihtiyac"], r["odenecek"], r["ertelenen"], r["kalan"], r["ileri_toplam"], r["ertelenen_gunluk"]),
                             (["Vergi Dairesi", "Yazılım Ltd."], D(1030000), D(875000), D(155000), D(25000), D(348000), D("40.00")))
            self.assertEqual([x["gunluk"] for x in r["sirali"][2:4]], [D("96.00"), D("150.00")])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Yüksek", "Nakit yetersiz", "Dönem"), ("Orta", "Ertelenen ödeme", "Ambalaj Ltd. · FT-0417"), ("Orta", "Ertelenen ödeme", "Lojistik A.Ş. · FT-7730")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Ödeme Sırası", "Özet", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "o.csv"
            y.write_text("Tedarikçi;Fatura No;Vade Tarihi;Kalan Tutar;Öncelik;Gecikme Cezası;Ertelenemez\nA;F1;14.10.2026;1.250,50;KRİTİK;2,5;x\nB;F2;14.10.2026;100;acil;-1;\nC;F3;x;100;;;\nD;F4;14.10.2026;0;;;\n;F5;14.10.2026;5;;;\n",
                         encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["cari"], x["belge"], x["tutar"], x["onem"], x["faiz"], x["zorunlu"]) for x in l], [u["tur"] for u in uy]),
                             ([("A", "F1", D("1250.50"), 0, D("2.5"), True), ("B", "F2", D(100), 2, D(0), False)], ["Önem tanınmadı"] + ["Satır okunamadı"] * 3))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--odemeler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--nakit", "1.000", "--asgari", "100", "--tarih", "12.10.2026", "--gun", "10"]), 0)
            self.assertEqual((main.main(arg), main.main(arg + ["--nakit", "x"]), main.main(arg + ["--nakit", "5", "--asgari", "x"]), main.main(arg + ["--nakit", "5", "--gun", "0"]), main.main(arg + ["--nakit", "5", "--tarih", "x"])),
                             (2, 2, 2, 2, 2))
            self.assertEqual(main.main(["--odemeler", str(Path(tmp) / "yok.csv"), "--nakit", "5"]), 1)
            for icerik in ("Alacaklı;Tutar\nA;1\n", "Alacaklı;Vade;Tutar\nA;x;1\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg + ["--nakit", "5"]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
