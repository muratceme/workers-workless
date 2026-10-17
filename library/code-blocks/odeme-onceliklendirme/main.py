"""
Ödeme Önceliklendirme — Gerisi Bende kod bloğu
Finans › Finans Uzmanı

Eldeki nakitle hangi ödemelerin yapılacağını, hangilerinin erteleneceğini açık bir kuralla sıralar.
  Dönem               = --tarih gününden başlayan --gun (7) gün;  vadesi dönem sonundan sonra olan kalem "Vadesi gelmedi" sayılır
  Kullanılabilir nakit= --nakit − --asgari (kasada bırakılacak tutar)
  Öncelik sırası      = 1) Zorunlu = Evet olanlar  2) Önem (Kritik → Yüksek → Normal → Düşük)  3) vadesi geçmiş, sonra dönem içi
                        4) günlük gecikme maliyeti yüksek olan  5) vadesi erken olan  6) tutarı küçük olan
  Günlük gecikme mal. = tutar × aylık gecikme faizi % ÷ 30
  Dağıtım             = sırayla; kalem kalan nakde sığıyorsa "Öde", sığmıyorsa "Ertele" denir ve sıradaki kaleme geçilir (kısmi ödeme önerilmez)
"Zorunlu" ve "Önem" sütunlarını siz doldurursunuz; kod hangi ödemenin yasal olarak ertelenemeyeceğine karar vermez.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
SUTUN = {"cari": ("alacakli", "cari", "tedarikci", "firma", "kime"), "belge": ("belge no", "belge", "fatura no", "aciklama"), "vade": ("vade", "vade tarihi"), "tutar": ("tutar", "kalan tutar", "odenecek tutar"),
         "tur": ("tur", "odeme turu"), "onem": ("onem", "oncelik", "tedarikci onemi"), "faiz": ("gecikme faizi", "gecikme faizi aylik", "aylik gecikme faizi", "gecikme cezasi"), "zorunlu": ("zorunlu", "ertelenemez")}
ONEM = {"kritik": 0, "cok yuksek": 0, "yuksek": 1, "onemli": 1, "": 2, "normal": 2, "orta": 2, "dusuk": 3}
ONEM_AD = ("Kritik", "Yüksek", "Normal", "Düşük")
Z, GUN = Decimal(0), timedelta(days=1)
EVET = ("evet", "e", "x", "1", "var", "true")
KARAR_RENK = {"Öde": YESIL, "Ertele": KIRMIZI, "Vadesi gelmedi": None}


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("cari", "vade", "tutar")):
        c = metin(r["cari"])
        if not c and not metin(r["tutar"]):
            continue
        kim = f"{c or '?'} · {metin(r['belge'])}".strip(" ·")
        v, t, f = tarih(r["vade"]), para(r["tutar"]), para(r["faiz"])
        if not c or v is None or t is None or t <= 0:
            uy.append(uyari("Yüksek", "Satır okunamadı", f"{kim} (satır {r['_satir']})", f"Vade '{metin(r['vade'])}', tutar '{metin(r['tutar'])}'. Tutar pozitif olmalı. Kalem listeye alınmadı."))
            continue
        o = katla(r["onem"])
        if o not in ONEM:
            uy.append(uyari("Orta", "Önem tanınmadı", kim, f"'{metin(r['onem'])}' tanınmadı; kalem 'Normal' sayıldı. Kritik, Yüksek, Normal veya Düşük yazın."))
        if f is not None and f < 0:
            f = None
        l.append({"cari": c, "belge": metin(r["belge"]), "vade": v, "tutar": t, "tur": metin(r["tur"]), "onem": ONEM.get(o, 2), "faiz": f or Z, "zorunlu": katla(r["zorunlu"]) in EVET, "kim": kim})
    return l, uy


def onceliklendir(kalemler: list[dict], nakit: Decimal, bas: date, gun: int = 7, asgari: Decimal = Z) -> dict:
    if not kalemler:
        raise ValueError("Sıralanacak ödeme kalemi yok.")
    son, uy = bas + (gun - 1) * GUN, []
    for k in kalemler:
        k["gecikme"] = max((bas - k["vade"]).days, 0)
        k["grup"] = "İleri vadeli" if k["vade"] > son else "Vadesi geçmiş" if k["gecikme"] else "Dönem içi"
        k["gunluk"] = yuvarla(k["tutar"] * k["faiz"] / 3000)
    sirali = sorted((k for k in kalemler if k["grup"] != "İleri vadeli"), key=lambda k: (not k["zorunlu"], k["onem"], k["grup"] != "Vadesi geçmiş", -k["gunluk"], k["vade"], k["tutar"], katla(k["cari"])))
    kalan = kullanilabilir = nakit - asgari
    for i, k in enumerate(sirali, 1):
        k["sira"] = i
        k["gerekce"] = " · ".join(x for x in ("Zorunlu" if k["zorunlu"] else "", ONEM_AD[k["onem"]], f"{k['gecikme']} gün gecikmiş" if k["gecikme"] else "dönem içi", f"günlük gecikme maliyeti {k['gunluk']:,.2f}" if k["gunluk"] else "") if x)
        if k["tutar"] <= kalan:
            k["karar"], kalan = "Öde", kalan - k["tutar"]
        else:
            k["karar"] = "Ertele"
            if k["zorunlu"] or k["onem"] == 0:
                uy.append(uyari("Yüksek", "Zorunlu ödeme karşılanamıyor" if k["zorunlu"] else "Kritik ödeme erteleniyor", k["kim"],
                                f"{k['tutar']:,.2f} tutarındaki ödeme (vade {k['vade']:%d.%m.%Y}) için kalan nakit {max(kalan, Z):,.2f}. Kaynak bulun veya alacaklıyla görüşün."))
            else:
                uy.append(uyari("Orta", "Ertelenen ödeme", k["kim"], f"{k['tutar']:,.2f} · vade {k['vade']:%d.%m.%Y}" + (f" · {k['gecikme']} gün gecikmiş" if k["gecikme"] else "")
                                + (f" · her gün yaklaşık {k['gunluk']:,.2f} gecikme maliyeti" if k["gunluk"] else "") + ". Alacaklıya bilgi verin."))
        k["kalan_nakit"] = kalan
    ileri = sorted((k for k in kalemler if k["grup"] == "İleri vadeli"), key=lambda k: (k["vade"], katla(k["cari"])))
    for k in ileri:
        k.update(sira=None, karar="Vadesi gelmedi", gerekce=("Zorunlu · " if k["zorunlu"] else "") + f"vadeye {(k['vade'] - bas).days} gün var", kalan_nakit=None)
    top = lambda karar: sum((k["tutar"] for k in sirali if k["karar"] == karar), Z)   # noqa: E731
    ihtiyac = sum((k["tutar"] for k in sirali), Z)
    if ihtiyac > kullanilabilir:
        uy.append(uyari("Yüksek", "Nakit yetersiz", "Dönem", f"Dönem ihtiyacı {ihtiyac:,.2f}, kullanılabilir nakit {max(kullanilabilir, Z):,.2f}; açık {ihtiyac - max(kullanilabilir, Z):,.2f}."))
    return {"sirali": sirali, "ileri": ileri, "bas": bas, "son": son, "nakit": nakit, "asgari": asgari, "kullanilabilir": kullanilabilir, "ihtiyac": ihtiyac, "odenecek": top("Öde"), "ertelenen": top("Ertele"),
            "kalan": kalan, "ileri_toplam": sum((k["tutar"] for k in ileri), Z), "ertelenen_gunluk": sum((k["gunluk"] for k in sirali if k["karar"] == "Ertele"), Z), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ödeme Sırası"
    _baslik(ws, ["Sıra", "Karar", "Alacaklı", "Belge", "Tür", "Vade", "Tutar", "Gerekçe", "Günlük Gecikme Maliyeti", "Ödeme Sonrası Kalan Nakit", "Onay / Not (siz doldurun)"], (6, 15, 28, 22, 13, 12, 16, 54, 14, 18, 28))
    for k in r["sirali"] + r["ileri"]:
        _satir(ws, [k["sira"], k["karar"], k["cari"], k["belge"], k["tur"], k["vade"], k["tutar"], k["gerekce"], k["gunluk"] or None, k["kalan_nakit"] if k["karar"] == "Öde" else None, ""], para_s=(7, 9, 10), tarih_s=(6,),
               dolgu={2: KARAR_RENK[k["karar"]]}, sar=(8,))
    ws = wb.create_sheet("Özet")
    _baslik(ws, ["Kalem", "Tutar"], (44, 18))
    for etiket, tutar in (("Eldeki nakit", r["nakit"]), ("Kasada bırakılacak (asgari)", r["asgari"]), ("Kullanılabilir nakit", r["kullanilabilir"]), ("Dönem ihtiyacı (vadesi geçmiş + dönem içi)", r["ihtiyac"]),
                          ("Ödenecek", r["odenecek"]), ("Ertelenen", r["ertelenen"]), ("Ödemeler sonrası kullanılabilir nakit", r["kalan"]), ("Ertelenenlerin günlük gecikme maliyeti", r["ertelenen_gunluk"]),
                          ("İleri vadeli (bu dönem ödenmeyecek)", r["ileri_toplam"])):
        _satir(ws, [etiket, tutar], para_s=(2,))
    ws.append([])
    ws.append([f"Dönem {r['bas']:%d.%m.%Y} – {r['son']:%d.%m.%Y}. Sıra: zorunlu → önem → vadesi geçmiş → günlük gecikme maliyeti → vade → tutar. Öneridir; ödeme talimatı değildir."])
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, nakit: Decimal, bas: date, gun: int = 7, asgari: Decimal = Z) -> dict:
    kalemler, uy = oku(yol)
    r = onceliklendir(kalemler, nakit, bas, gun, asgari)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Eldeki nakde göre ödeme listesini önceliklendirir.")
    p.add_argument("--odemeler", type=Path, help="Alacaklı, Belge No, Vade, Tutar, Tür, Önem, Gecikme Faizi (aylık %%), Zorunlu")
    p.add_argument("--nakit", help="Eldeki nakit (kasa ve banka)")
    p.add_argument("--asgari", default="0", help="Kasada bırakılacak asgari tutar (varsayılan 0)")
    p.add_argument("--tarih", help="Dönem başlangıcı (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--gun", type=int, default=7, help="Dönem uzunluğu, gün (varsayılan 7)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "odeme_onceliklendirme.xlsx")
    a = p.parse_args(argv)
    ornek = a.odemeler is None
    yol = ORNEK / "odemeler.csv" if ornek else a.odemeler
    bas = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    nakit, asgari = para(a.nakit) if a.nakit else Decimal(900000) if ornek else None, para(a.asgari)
    if nakit is None or nakit < 0 or asgari is None or asgari < 0 or bas is None or not 1 <= a.gun <= 92:
        print("[X] --nakit zorunludur ve sıfır veya pozitif olmalı; --asgari sayı, --tarih GG.AA.YYYY, --gun 1–92 arasında olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, nakit, bas, a.gun, asgari)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {r['bas']:%d.%m.%Y} – {r['son']:%d.%m.%Y} · ihtiyaç {r['ihtiyac']:,.2f} · ödenecek {r['odenecek']:,.2f} · ertelenen {r['ertelenen']:,.2f} · kalan nakit {r['kalan']:,.2f}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
