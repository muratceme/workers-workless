import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
B = date(2026, 10, 12)
ek = lambda g: date(2026, 10, g)   # noqa: E731


def k(bas, bit, kontenjan, release, acente="A", oda="STD"):
    return {"acente": acente, "oda": oda, "bas": ek(bas), "bit": ek(bit), "kontenjan": kontenjan, "release": release}


def rz(giris, cikis, adet, acente="A", oda="STD"):
    return {"acente": acente, "oda": oda, "giris": ek(giris), "cikis": ek(cikis), "adet": adet}


REZ = lambda: [rz(14, 17, 4), rz(15, 16, 2), rz(19, 21, 5), rz(8, 13, 1)]   # noqa: E731


class TakipTesti(unittest.TestCase):
    def test_gunluk_durum(self):
        r = main.takip([k(10, 20, 5, 3)], REZ(), B)
        self.assertEqual([(g["gun"].day, g["kullanilan"], g["kalan"], g["release_tarihi"].day, g["kalan_gun"], g["durum"]) for g in r["gunler"]],
                         [(12, 1, 4, 9, -3, "Release oldu"), (13, 0, 5, 10, -2, "Release oldu"), (14, 4, 1, 11, -1, "Release oldu"), (15, 6, -1, 12, 0, "Aşım"), (16, 4, 1, 13, 1, "Release yaklaşıyor"),
                          (17, 0, 5, 14, 2, "Release yaklaşıyor"), (18, 0, 5, 15, 3, "Release yaklaşıyor"), (19, 5, 0, 16, 4, "Dolu"), (20, 5, 0, 17, 5, "Dolu")])
        o = r["ozet"][0]
        self.assertEqual((o["gun"], o["kontenjan"], o["kullanilan"], o["kullanim"], o["iade"], o["yaklasan"], o["asim"], o["satilabilir"], o["asim_gunleri"]), (9, 45, 25, D("55.56"), 10, 11, 1, 11, [ek(15)]))
        self.assertEqual((r["say"], [g["gun"].day for g in r["yaklasan"]], sorted(u["tur"] for u in r["uyarilar"])),
                         ({"Release oldu": 3, "Aşım": 1, "Release yaklaşıyor": 3, "Dolu": 2}, [16, 17, 18], ["Kontenjan aşımı", "Release yaklaşıyor"]))

    def test_ayarlar(self):
        r = main.takip([k(10, 20, 5, 3)], REZ(), B, yaklasan=1)
        self.assertEqual(([g["durum"] for g in r["gunler"]][4:7], r["ozet"][0]["yaklasan"], r["ozet"][0]["satilabilir"]), (["Release yaklaşıyor", "Açık", "Açık"], 1, 11))
        r = main.takip([k(10, 20, 5, 3)], REZ(), B, ufuk=3)
        self.assertEqual(([g["gun"].day for g in r["gunler"]], r["son"], r["ozet"][0]["kullanilan"]), ([12, 13, 14], ek(14), 5))
        r = main.takip([k(10, 20, 5, 0)], [], B, ufuk=2)
        self.assertEqual([(g["durum"], g["kalan_gun"]) for g in r["gunler"]], [("Release oldu", 0), ("Release yaklaşıyor", 1)])           # release 0: yalnız bugünün kontenjanı dönmüş
        for kont, bugun in (([k(1, 11, 5, 3)], B), ([k(10, 20, 5, 3)], date(2026, 9, 1)), ([], B)):
            with self.assertRaises(ValueError):
                main.takip(kont, [], bugun)

    def test_birden_cok_kontrat(self):
        r = main.takip([k(12, 13, 5, 3), k(13, 14, 2, 7), k(12, 12, 4, 1, oda="Aile"), k(12, 12, 3, 1, acente="B")], [rz(12, 14, 6), rz(12, 13, 1, oda="aile"), rz(12, 13, 2, oda="Suit"), rz(12, 14, 1, acente="C")], B)
        self.assertEqual([(g["acente"], g["oda"], g["gun"].day, g["kontenjan"], g["kullanilan"], g["durum"]) for g in r["gunler"]],
                         [("A", "Aile", 12, 4, 1, "Release oldu"), ("A", "STD", 12, 5, 6, "Aşım"), ("A", "STD", 13, 7, 6, "Release oldu"), ("A", "STD", 14, 2, 0, "Release oldu"), ("B", "STD", 12, 3, 0, "Release oldu")])
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]), [("Kontenjan aşımı", "A · STD"), ("Kontenjan dışı rezervasyon", "A · Suit"), ("Kontenjan dışı rezervasyon", "C · STD"), ("Çakışan kontrat satırları", "A · STD")])
        self.assertIn("2 oda-gece", [u for u in r["uyarilar"] if u["kim"] == "C · STD"][0]["aciklama"])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "k.xlsx"
            r = main.calistir(ORNEK / "kontratlar.csv", ORNEK / "rezervasyonlar.csv", cikti, B)
            self.assertEqual([(o["acente"], o["oda"], o["gun"], o["kontenjan"], o["kullanilan"], o["kullanim"], o["iade"], o["yaklasan"], o["satilabilir"], o["asim"]) for o in r["ozet"]],
                             [("Deneme Travel", "Standart", 27, 162, 15, D("9.26"), 72, 27, 75, 0), ("Örnek Tur", "Aile", 20, 80, 10, D("12.50"), 32, 18, 38, 0), ("Örnek Tur", "Standart", 20, 200, 44, D("22.00"), 38, 70, 120, 2)])
            self.assertEqual((len(r["gunler"]), r["son"], [(u["tur"], u["kim"]) for u in r["uyarilar"]]),
                             (67, date(2026, 11, 10), [("Kontenjan aşımı", "Örnek Tur · Standart"), ("Release yaklaşıyor", "Deneme Travel · Standart"), ("Release yaklaşıyor", "Örnek Tur · Aile"),
                                                      ("Release yaklaşıyor", "Örnek Tur · Standart"), ("Kontenjan dışı rezervasyon", "Örnek Tur · Suit")]))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Günlük Durum", "Yaklaşan Release", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            ky, ry = Path(tmp) / "k.csv", Path(tmp) / "r.csv"
            ky.write_text("Tur Operatörü;Oda Tipi;Başlangıç;Bitiş;Allotment;Release\nA;STD;01.10.2026;31.10.2026;5;\nA;STD;31.10.2026;01.10.2026;5;3\nA;STD;01.10.2026;31.10.2026;2,5;3\nA;;01.10.2026;31.10.2026;5;3\n", encoding="utf-8")
            ry.write_text("Acente;Oda Tipi;Giriş Tarihi;Çıkış Tarihi;Adet;Durum\nA;STD;14.10.2026;16.10.2026;;\nA;STD;14.10.2026;14.10.2026;1;\nA;STD;14.10.2026;16.10.2026;3;İPTAL EDİLDİ\nA;STD;x;16.10.2026;1;\n", encoding="utf-8")
            kont, rez, uy = main.oku(ky, ry)
            self.assertEqual((kont, rez, sorted(u["tur"] for u in uy)), ([{"acente": "A", "oda": "STD", "bas": ek(1), "bit": ek(31), "kontenjan": 5, "release": 0}], [{"acente": "A", "oda": "STD", "giris": ek(14), "cikis": ek(16), "adet": 1}],
                                                                     ["Kontrat satırı okunamadı"] * 3 + ["Rezervasyon okunamadı"] * 2))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--kontratlar", str(ky), "--cikti", c, "--tarih", "12.10.2026"]
            self.assertEqual(main.main(arg + ["--rezervasyonlar", str(ry), "--ufuk", "10", "--yaklasan", "3"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(["--kontratlar", str(ky), "--cikti", c, "--tarih", "01.12.2026"]), 1)
            for ekarg in (["--ufuk", "0"], ["--yaklasan", "-1"], ["--tarih", "bugün"]):
                self.assertEqual(main.main(arg + ekarg), 2)
            self.assertEqual(main.main(arg + ["--rezervasyonlar", str(Path(tmp) / "yok.csv")]), 1)
            ky.write_text("Acente;Oda Tipi;Başlangıç;Bitiş;Kontenjan\nA;STD;x;y;5\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
