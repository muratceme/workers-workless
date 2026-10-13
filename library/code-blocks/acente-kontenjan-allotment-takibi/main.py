"""
Acente Kontenjan (Allotment) Takibi — Workers / Workless kod bloğu
Rezervasyon ve Gelir Yönetimi › Rezervasyon Sorumlusu

Tur operatörü ve acente kontratlarındaki kontenjanı, release tarihini ve kullanımını konaklama günü bazında izler.
  Kullanılan        = o gece konaklayan (giriş ≤ gün < çıkış) ve iptal olmayan rezervasyonların oda sayısı toplamı
  Kalan             = kontenjan − kullanılan
  Release tarihi    = konaklama günü − release gün sayısı;  bugün ≥ release tarihi ise kullanılmayan kontenjan otele döner ("Release oldu")
  Durum             = Aşım (kullanılan > kontenjan) · Release oldu · Dolu · Release yaklaşıyor (release'e --yaklasan gün veya daha az, kalan > 0) · Açık
  Özet              = ufuktaki oda-gece: kontenjan, kullanılan, kullanım %, release ile dönen, release'i yaklaşan boş, aşım
Yalnız bugünden itibaren --ufuk gün (varsayılan 30) içindeki konaklama günleri izlenir. Release tanımı (gün sayımı, saat) kontrata göre değişir; kendi kontratınızı esas alın.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


KONTRAT = {"acente": ("acente", "acenta", "tur operatoru", "operator"), "oda": ("oda tipi", "oda tipi kodu", "tip"), "bas": ("baslangic", "baslangic tarihi", "kontrat baslangic"),
           "bit": ("bitis", "bitis tarihi", "kontrat bitis"), "kontenjan": ("kontenjan", "kontenjan oda", "gunluk kontenjan", "allotment"), "release": ("release", "release gun", "release suresi", "release gunu")}
REZ = {"acente": KONTRAT["acente"], "oda": KONTRAT["oda"], "giris": ("giris", "giris tarihi", "check in"), "cikis": ("cikis", "cikis tarihi", "check out"),
       "adet": ("oda sayisi", "adet", "oda adedi"), "durum": ("durum", "rezervasyon durumu")}
GUN = timedelta(days=1)
DURUM_RENK = {"Aşım": KIRMIZI, "Release oldu": MAVI, "Dolu": YESIL, "Release yaklaşıyor": SARI}


def tam(x) -> int | None:
    v = para(x)
    return int(v) if v is not None and v == v.to_integral_value() else None


def oku(kontrat_yol: Path, rez_yol: Path | None = None) -> tuple[list[dict], list[dict], list[dict]]:
    uy, kontratlar, rezler = [], [], []
    for r in kayitlar(kontrat_yol, KONTRAT, ("acente", "oda", "bas", "bit", "kontenjan")):
        a, o = metin(r["acente"]), metin(r["oda"])
        if not a and not o:
            continue
        b, s, k, rl = tarih(r["bas"]), tarih(r["bit"]), tam(r["kontenjan"]), tam(r["release"]) if metin(r["release"]) else 0
        if not a or not o or b is None or s is None or s < b or k is None or k <= 0 or rl is None or rl < 0:
            uy.append(uyari("Yüksek", "Kontrat satırı okunamadı", f"{a or '?'} · {o or '?'} (satır {r['_satir']})", f"Başlangıç '{metin(r['bas'])}', bitiş '{metin(r['bit'])}', kontenjan '{metin(r['kontenjan'])}', release '{metin(r['release'])}'. "
                            "Kontenjan pozitif tam sayı, release sıfır veya pozitif tam sayı olmalı. Satır alınmadı."))
            continue
        kontratlar.append({"acente": a, "oda": o, "bas": b, "bit": s, "kontenjan": k, "release": rl})
    if rez_yol:
        for r in kayitlar(rez_yol, REZ, ("acente", "oda", "giris", "cikis")):
            a, o = metin(r["acente"]), metin(r["oda"])
            if not a and not o:
                continue
            if katla(r["durum"]).startswith(("iptal", "cancel")):
                continue
            g, c, n = tarih(r["giris"]), tarih(r["cikis"]), tam(r["adet"]) if metin(r["adet"]) else 1
            if not a or not o or g is None or c is None or c <= g or n is None or n <= 0:
                uy.append(uyari("Yüksek", "Rezervasyon okunamadı", f"{a or '?'} · {o or '?'} (satır {r['_satir']})", f"Giriş '{metin(r['giris'])}', çıkış '{metin(r['cikis'])}', oda sayısı '{metin(r['adet'])}'. Çıkış girişten sonra olmalı. Satır alınmadı; kullanım eksik sayıldı."))
                continue
            rezler.append({"acente": a, "oda": o, "giris": g, "cikis": c, "adet": n})
    return kontratlar, rezler, uy


def takip(kontratlar: list[dict], rezler: list[dict], bugun: date, ufuk: int = 30, yaklasan: int = 7) -> dict:
    son, uy, kont, cakisan = bugun + timedelta(days=ufuk - 1), [], {}, set()
    for k in kontratlar:
        g = max(k["bas"], bugun)
        while g <= min(k["bit"], son):
            a = (katla(k["acente"]), katla(k["oda"]), g)
            if a in kont:
                kont[a]["kontenjan"] += k["kontenjan"]
                kont[a]["release"] = max(kont[a]["release"], k["release"])
                cakisan.add((k["acente"], k["oda"]))
            else:
                kont[a] = {"acente": k["acente"], "oda": k["oda"], "gun": g, "kontenjan": k["kontenjan"], "release": k["release"]}
            g += GUN
    if not kont:
        raise ValueError(f"{bugun:%d.%m.%Y} – {son:%d.%m.%Y} aralığında geçerli kontenjan yok.")
    kul, disi = defaultdict(int), defaultdict(int)
    for r in rezler:
        g = max(r["giris"], bugun)
        while g < r["cikis"] and g <= son:
            a = (katla(r["acente"]), katla(r["oda"]), g)
            if a in kont:
                kul[a] += r["adet"]
            else:
                disi[(r["acente"], r["oda"])] += r["adet"]
            g += GUN
    gunler, ozet = [], {}
    for a in sorted(kont, key=lambda a: (kont[a]["acente"], kont[a]["oda"], a[2])):
        k = kont[a]
        kullanilan, rt = kul[a], k["gun"] - timedelta(days=k["release"])
        kalan, kg = k["kontenjan"] - kullanilan, (rt - bugun).days
        durum = "Aşım" if kalan < 0 else "Release oldu" if kg <= 0 else "Dolu" if kalan == 0 else "Release yaklaşıyor" if kg <= yaklasan else "Açık"
        gunler.append(k | {"kullanilan": kullanilan, "kalan": kalan, "release_tarihi": rt, "kalan_gun": kg, "durum": durum})
        o = ozet.setdefault(a[:2], {"acente": k["acente"], "oda": k["oda"], "gun": 0, "kontenjan": 0, "kullanilan": 0, "iade": 0, "yaklasan": 0, "asim": 0, "satilabilir": 0, "asim_gunleri": []})
        o["gun"] += 1
        o["kontenjan"] += k["kontenjan"]
        o["kullanilan"] += kullanilan
        if durum == "Aşım":
            o["asim"] -= kalan
            o["asim_gunleri"].append(k["gun"])
        elif durum == "Release oldu":
            o["iade"] += kalan
        elif durum == "Release yaklaşıyor":
            o["yaklasan"] += kalan
        if durum in ("Release yaklaşıyor", "Açık"):
            o["satilabilir"] += kalan
    for o in ozet.values():
        o["kullanim"] = yuvarla(Decimal(o["kullanilan"]) / o["kontenjan"] * 100)
        if o["asim"]:
            g = o["asim_gunleri"]
            uy.append(uyari("Yüksek", "Kontenjan aşımı", f"{o['acente']} · {o['oda']}", f"{len(g)} gecede toplam {o['asim']} oda-gece kontenjan üstü rezervasyon var ({', '.join(f'{x:%d.%m}' for x in g[:8])}{'…' if len(g) > 8 else ''}). "
                            "Ek kontenjan onayı var mı, kontrol edin."))
        if o["yaklasan"]:
            uy.append(uyari("Orta", "Release yaklaşıyor", f"{o['acente']} · {o['oda']}", f"Release'ine {yaklasan} gün veya daha az kalan günlerde {o['yaklasan']} oda-gece boş kontenjan var. Acenteden satış durumunu isteyin."))
    for (a, o), n in sorted(disi.items()):
        uy.append(uyari("Bilgi", "Kontenjan dışı rezervasyon", f"{a} · {o}", f"Ufuktaki {n} oda-gece için bu acente ve oda tipinde geçerli kontenjan yok (kontrat dışı veya serbest satış)."))
    for a, o in sorted(cakisan):
        uy.append(uyari("Bilgi", "Çakışan kontrat satırları", f"{a} · {o}", "Aynı günleri kapsayan birden fazla kontrat satırı var. Kontenjanlar toplandı, release için en uzun süre alındı."))
    yak = sorted((g for g in gunler if g["durum"] == "Release yaklaşıyor"), key=lambda g: (g["release_tarihi"], g["acente"], g["oda"]))
    say = defaultdict(int)
    for g in gunler:
        say[g["durum"]] += 1
    return {"gunler": gunler, "ozet": list(ozet.values()), "yaklasan": yak, "say": dict(say), "bugun": bugun, "son": son, "ayar": (ufuk, yaklasan), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    _baslik(ws, ["Acente", "Oda Tipi", "Gün Sayısı", "Kontenjan (oda-gece)", "Kullanılan", "Kullanım %", "Release ile Dönen", "Release'i Yaklaşan Boş", "Satılabilir Kalan", "Aşım"], (28, 16, 9, 12, 11, 10, 12, 13, 12, 8))
    for o in r["ozet"]:
        _satir(ws, [o["acente"], o["oda"], o["gun"], o["kontenjan"], o["kullanilan"], o["kullanim"], o["iade"] or None, o["yaklasan"] or None, o["satilabilir"] or None, o["asim"] or None], para_s=(6,),
               dolgu={10: KIRMIZI if o["asim"] else None, 8: SARI if o["yaklasan"] else None})
    ws.append([])
    ws.append([f"Kapsam: {r['bugun']:%d.%m.%Y} – {r['son']:%d.%m.%Y} ({r['ayar'][0]} gün). Release tarihi = konaklama günü − release gün sayısı; bugün bu tarihe ulaştıysa boş kontenjan otele dönmüş sayılır. Kontratınızdaki tanımı esas alın."])
    for ad, satirlar in (("Günlük Durum", r["gunler"]), ("Yaklaşan Release", r["yaklasan"])):
        ws = wb.create_sheet(ad)
        _baslik(ws, ["Konaklama Günü", "Acente", "Oda Tipi", "Kontenjan", "Kullanılan", "Kalan", "Release Tarihi", "Release'e Kalan Gün", "Durum"], (14, 28, 16, 10, 10, 8, 13, 11, 20))
        for g in satirlar:
            _satir(ws, [g["gun"], g["acente"], g["oda"], g["kontenjan"], g["kullanilan"], g["kalan"], g["release_tarihi"], max(g["kalan_gun"], 0), g["durum"]], tarih_s=(1, 7), dolgu={9: DURUM_RENK.get(g["durum"])})
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(kontrat_yol: Path, rez_yol: Path | None, cikti: Path, bugun: date, **ayar) -> dict:
    kontratlar, rezler, uy = oku(kontrat_yol, rez_yol)
    if not kontratlar:
        raise ValueError("Kontrat dosyasında geçerli satır yok.")
    r = takip(kontratlar, rezler, bugun, **ayar)
    uy += r["uyarilar"]
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Acente kontenjanlarını, release tarihlerini ve kullanımı konaklama günü bazında izler.")
    p.add_argument("--kontratlar", type=Path, default=ORNEK / "kontratlar.csv", help="Acente, Oda Tipi, Başlangıç, Bitiş, Kontenjan, Release (gün)")
    p.add_argument("--rezervasyonlar", type=Path, help="Acente, Oda Tipi, Giriş, Çıkış, Oda Sayısı, Durum")
    p.add_argument("--tarih", help="Bugünün tarihi (GG.AA.YYYY); verilmezse sistem tarihi")
    p.add_argument("--ufuk", type=int, default=30, help="İzlenecek gün sayısı (varsayılan 30)")
    p.add_argument("--yaklasan", type=int, default=7, help="Release'e bu kadar gün veya daha az kalan günler işaretlenir (varsayılan 7)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kontenjan_takibi.xlsx")
    a = p.parse_args(argv)
    ornek = a.kontratlar == ORNEK / "kontratlar.csv"
    rez = a.rezervasyonlar or (ORNEK / "rezervasyonlar.csv" if ornek else None)
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    if bugun is None or not 1 <= a.ufuk <= 730 or a.yaklasan < 0:
        print("[X] --tarih GG.AA.YYYY biçiminde, --ufuk 1–730 arasında, --yaklasan sıfır veya pozitif olmalı")
        return 2
    for y in (a.kontratlar, rez):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.kontratlar, rez, a.cikti, bugun, ufuk=a.ufuk, yaklasan=a.yaklasan)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {r['bugun']:%d.%m.%Y} – {r['son']:%d.%m.%Y} · {len(r['ozet'])} acente-oda tipi · {len(r['gunler'])} gün satırı")
    for o in r["ozet"]:
        print(f"[{'!' if o['asim'] else 'OK'}] {o['acente']} · {o['oda']}: kontenjan {o['kontenjan']} oda-gece, kullanım %{o['kullanim']}, release ile dönen {o['iade']}, release'i yaklaşan boş {o['yaklasan']}"
              + (f", aşım {o['asim']}" if o["asim"] else ""))
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
