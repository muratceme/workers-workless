import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 8)


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "k.xlsx"
        cls.s = main.calistir(ORNEK / "kumas_siparisleri.csv", ORNEK / "asama_sureleri.csv", cls.cikti, ORNEK / "siparisler.csv", BUGUN)
        cls.k = {k.no: k for k in cls.s["kumaslar"]}
        cls.e = {e["po"]: e for e in cls.s["etkiler"]}
        cls.u = {}
        for u in cls.s["uyarilar"]:
            cls.u.setdefault(u["tur"], []).append(u)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_tahmini_teslim(self):
        self.assertEqual({k: (x.hazir, x.fark, x.durum) for k, x in self.k.items()},
                         {"KS-101": (date(2026, 10, 3), -2, "Teslim alındı"), "KS-102": (date(2026, 10, 15), 9, "TERMİN GEÇTİ"),
                          "KS-103": (date(2026, 10, 12), 0, "Zamanında"), "KS-104": (date(2026, 10, 20), 5, "Gecikecek"),
                          "KS-105": (date(2026, 10, 6), 2, "Eksik teslim"), "KS-106": (date(2026, 10, 27), 2, "Gecikecek")})
        self.assertEqual(self.k["KS-102"].tahmini["Boyama"], BUGUN)                      # gecikmiş aşama bugün biter
        self.assertEqual(self.k["KS-104"].tahmini["Boyama"], date(2026, 10, 13))         # lab dip bekliyor: bugün başlar
        self.assertEqual(self.k["KS-104"].mevcut_asama, "Boyama (lab dip onayı bekleniyor)")
        self.assertNotIn("Örme", self.k["KS-105"].asamalar[0])

    def test_siparis_etkisi(self):
        e = self.e["PO-501"]
        self.assertEqual((e["kritik"].no, e["kayma"], e["tahmini_sevk"], e["durum"]), ("KS-102", 2, date(2026, 11, 22), "Kesim kayıyor"))
        self.assertEqual(self.e["PO-502"]["durum"], "Riskli")
        self.assertEqual(self.e["PO-503"]["durum"], "Zamanında · Eksik kumaş")
        self.assertEqual(self.e["PO-504"]["kayma"], 0)

    def test_uyarilar(self):
        self.assertEqual(self.s["uyarilar"][0]["tur"], "Kesim / sevk kayıyor")
        self.assertIn("%7,1 eksik", self.u["Eksik teslim"][0]["aciklama"])
        self.assertEqual(self.u["Lab dip onayı yok"][0]["kumas"], "KS-104")
        self.assertEqual(sorted(u["kumas"] for u in self.u["Aşama gecikti"]), ["KS-102", "KS-103", "KS-106"])
        self.assertIn("bolluk yok", self.u["Kesim için az bolluk"][0]["aciklama"])
        self.assertNotIn("Aşama tarihi boş", self.u)

    def test_performans(self):
        p = {x["ad"]: x for x in self.s["performans"]}
        self.assertEqual((p["Tedarikçi A Örme"]["zamaninda"], p["Tedarikçi A Örme"]["acik_gec"]), (1, 2))
        self.assertEqual(p["Tedarikçi C Dokuma"]["ort_gecikme"], 2)

    def test_excel(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Sipariş Etkisi", "Kumaş Durumu", "Aşama Matrisi", "Tedarikçi Performansı", "Uyarılar"])
        am = wb["Aşama Matrisi"]
        basliklar = [c.value for c in am[1]]
        self.assertEqual(basliklar[:5], ["Kumaş Sipariş No", "Kumaş", "Lab Dip Onayı", "İplik Temini", "Örme"])
        self.assertIn("Dokuma", basliklar)
        self.assertEqual(wb["Sipariş Etkisi"]["I2"].value, 2)


class KuralTesti(unittest.TestCase):
    def test_sira_hatasi_ve_bos_ara_asama(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t / "s.csv").write_text("Aşama;Süre (gün)\nÖrme;5\nBoyama;4\nSevk;1\n", encoding="utf-8")
            (t / "k.csv").write_text("Kumaş Sipariş No;Termin;Sipariş Tarihi;Örme;Boyama;Sevk\nA;20.10.2026;01.10.2026;06.10.2026;05.10.2026;\n"
                                     "B;20.10.2026;01.10.2026;;07.10.2026;\n", encoding="utf-8")
            s = main.calistir(t / "k.csv", t / "s.csv", t / "o.xlsx", bugun=BUGUN)
            turler = {(u["tur"], u["kumas"]) for u in s["uyarilar"]}
            self.assertIn(("Veri hatası", "A"), turler)
            self.assertIn(("Aşama tarihi boş", "B"), turler)
            self.assertNotIn("Lab dip onayı yok", {u["tur"] for u in s["uyarilar"]})     # sütun yok
            self.assertEqual(s["kumaslar"][1].tahmini_teslim, date(2026, 10, 8))

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--kumaslar", str(Path(tmp) / "yok.csv")]), 1)
            self.assertEqual(main.main(["--bugun", "8 Ekim", "--cikti", str(Path(tmp) / "x.xlsx")]), 1)


if __name__ == "__main__":
    unittest.main()
