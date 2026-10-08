"""
Kumaş Sipariş Termin Takibi — Workers / Workless kod bloğu
Tekstil ve Konfeksiyon › Kumaş ve Aksesuar Satın Alma › Kumaş Satın Alma Uzmanı

Kumaş siparişlerinin üretim aşamalarını (iplik, örme / dokuma, boyahane, apre, kalite kontrol, sevk) ve termin
durumunu takip eder; gecikmenin konfeksiyon siparişinin kesim ve sevk tarihine etkisini hesaplar:
  - Tahmini teslim: son tamamlanan aşamadan itibaren kalan aşamaların standart süreleri (takvim günü). Süresi
    geçmiş ama tamamlanmamış aşamanın en erken bugün biteceği varsayılır.
  - Termin karşılaştırması: teslim edilmiş siparişte gerçek gecikme, açık siparişte tahmini gecikme.
  - Lab dip onayı olmadan boyamaya sıra gelmiş siparişler.
  - Eksik / fazla teslim (tolerans varsayılan %3).
  - Sipariş etkisi: bağlı konfeksiyon siparişinin tüm kumaşları hazır olmadan kesim başlayamaz. Kumaş hazır
    + giriş kontrolü (varsayılan 2 gün) planlanan kesimi geçiyorsa kesim ve sevk aynı gün sayısı kadar kayar.
Rapor: kumaş durumu, aşama matrisi, sipariş etkisi, tedarikçi performansı, uyarılar. İnternete bağlanmaz.

Kullanım:
    python main.py                                                     # örnek: 6 kumaş siparişi, durum tarihi 08.10.2026
    python main.py --kumaslar kumas_siparisleri.xlsx --sureler asama_sureleri.csv --siparisler siparisler.xlsx --bugun 08.10.2026
"""
from __future__ import annotations

import argparse
import csv
import re
import statistics
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")

KUMAS_SUTUNLARI = {"no": ("kumas siparis no", "siparis no", "kumas po", "no"), "tedarikci": ("tedarikci", "firma", "kumasci"),
                   "kumas": ("kumas", "kalite", "kumas kalitesi"), "renk": ("renk",), "tip": ("tip", "kumas tipi", "uretim tipi"),
                   "miktar": ("miktar", "siparis miktari"), "birim": ("birim",), "po": ("bagli siparis", "konfeksiyon siparisi", "po", "siparis"),
                   "siparis": ("siparis tarihi", "kumas siparis tarihi"), "termin": ("termin", "termin tarihi", "soz verilen teslim"),
                   "labdip": ("lab dip onayi", "lab dip", "renk onayi"), "teslim_miktar": ("teslim alinan miktar", "teslim miktari", "gelen miktar"),
                   "teslim": ("teslim tarihi", "teslim alinan tarih", "giris tarihi")}
SURE_SUTUNLARI = {"tip": ("tip", "kumas tipi"), "asama": ("asama",), "sure": ("sure gun", "sure", "gun")}
SIPARIS_SUTUNLARI = {"no": ("siparis no", "po", "po no"), "model": ("model", "style"), "musteri": ("musteri", "buyer"),
                     "kesim": ("planlanan kesim tarihi", "kesim tarihi", "kesim"), "sevk": ("sevk tarihi", "ex factory", "termin")}
BOYAMA = ("boyama", "boyahane", "boya")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def sayi(x) -> float | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip().replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def binlik(x: float) -> str:
    return f"{x:,.0f}".replace(",", ".")


def oran_yaz(x: float) -> str:
    return f"{x * 100:.1f}".replace(".", ",").removesuffix(",0")


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...], ek_sutunlar: tuple[str, ...] = ()) -> list[dict]:
    """ek_sutunlar: başlığı katlanmış hâliyle birebir eşleşen ek sütunlar (aşama adları), '_ek' sözlüğüne konur."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            ek = {a: b.index(katla(a)) for a in ek_sutunlar if katla(a) in b}
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n, "_ek": {a: (r[j] if j < len(r) else None) for a, j in ek.items()}}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Kumas:
    satir: int
    no: str
    tedarikci: str
    kumas: str
    renk: str
    tip: str
    miktar: float | None
    birim: str
    po: str
    siparis: date | None
    termin: date | None
    labdip: date | None
    labdip_sutunu: bool
    teslim_miktar: float | None
    teslim: date | None
    gerceklesen: dict                       # aşama → tarih
    asamalar: list = field(default_factory=list)   # [(aşama, süre)]
    tahmini: dict = field(default_factory=dict)    # aşama → tahmini bitiş
    mevcut_asama: str = ""
    tahmini_teslim: date | None = None
    fark: int | None = None                  # + gecikme gün
    durum: str = ""

    @property
    def ad(self) -> str:
        return f"{self.no} {self.kumas} {self.renk}".strip()

    @property
    def hazir(self) -> date | None:
        return self.teslim or self.tahmini_teslim


def sureleri_oku(yol: Path) -> dict[str, list[tuple[str, int]]]:
    sonuc = defaultdict(list)
    for r in kayitlar(yol, SURE_SUTUNLARI, ("asama", "sure")):
        if r.get("asama") and sayi(r.get("sure")) is not None:
            sonuc[katla(r.get("tip")) or "*"].append((metin(r["asama"]), int(sayi(r["sure"]))))
    if not sonuc:
        raise ValueError(f"{yol.name}: aşama tanımı yok")
    return dict(sonuc)


def kumaslari_oku(yol: Path, sureler: dict) -> list[Kumas]:
    tum_asamalar = tuple(dict.fromkeys(a for lst in sureler.values() for a, _ in lst))
    sonuc = []
    for r in kayitlar(yol, KUMAS_SUTUNLARI, ("no", "termin"), tum_asamalar):
        if not r.get("no"):
            continue
        tip = katla(r.get("tip"))
        asamalar = sureler.get(tip) or sureler.get("*") or next(iter(sureler.values()))
        sonuc.append(Kumas(r["_satir"], metin(r["no"]), metin(r.get("tedarikci")), metin(r.get("kumas")), metin(r.get("renk")), metin(r.get("tip")),
                           sayi(r.get("miktar")), metin(r.get("birim")), metin(r.get("po")), tarih(r.get("siparis")), tarih(r.get("termin")),
                           tarih(r.get("labdip")), "labdip" in r, sayi(r.get("teslim_miktar")), tarih(r.get("teslim")),
                           {a: tarih(v) for a, v in r["_ek"].items() if tarih(v)}, list(asamalar)))
    return sonuc


def siparisleri_oku(yol: Path | None) -> dict[str, dict]:
    if yol is None:
        return {}
    return {metin(r["no"]): {"no": metin(r["no"]), "model": metin(r.get("model")), "musteri": metin(r.get("musteri")), "kesim": tarih(r.get("kesim")),
                             "sevk": tarih(r.get("sevk"))}
            for r in kayitlar(yol, SIPARIS_SUTUNLARI, ("no",)) if r.get("no")}


# ----------------------------------------------------------------------------
# Analiz
# ----------------------------------------------------------------------------

SIRA = {"Kritik": 0, "Yüksek": 1, "Orta": 2, "Bilgi": 3}


def analiz_et(kumaslar: list[Kumas], siparisler: dict, bugun: date, kontrol_gun: int = 2, tolerans: float = 0.03, risk_gun: int = 2) -> dict:
    uyarilar = []

    def uyar(onem, tur, k: Kumas | None, aciklama, po=""):
        uyarilar.append({"onem": onem, "tur": tur, "kumas": k.no if k else "", "po": k.po if k else po, "tedarikci": k.tedarikci if k else "",
                         "aciklama": aciklama})

    for k in kumaslar:
        adlar = [a for a, _ in k.asamalar]
        # Aşama tarih sırası
        onceki = None
        for a in adlar:
            t = k.gerceklesen.get(a)
            if t and onceki and t < onceki[1]:
                uyar("Orta", "Veri hatası", k, f"{a} ({t:%d.%m.%Y}) bir önceki aşamadan ({onceki[0]} {onceki[1]:%d.%m.%Y}) önce görünüyor")
            if t:
                onceki = (a, t)
        son_i = max((i for i, a in enumerate(adlar) if a in k.gerceklesen), default=-1)
        atlanan = [a for a in adlar[:max(son_i, 0)] if a not in k.gerceklesen]
        if atlanan:
            uyar("Bilgi", "Aşama tarihi boş", k, f"Sonraki aşama tamamlanmış ama şu aşamaların tarihi boş: {', '.join(atlanan)}")
        if k.teslim:
            k.durum, k.mevcut_asama = "Teslim alındı", "Teslim alındı"
            k.fark = (k.teslim - k.termin).days if k.termin else None
            if k.fark and k.fark > 0:
                uyar("Bilgi", "Geç teslim", k, f"{k.teslim:%d.%m.%Y} teslim alındı, termin {k.termin:%d.%m.%Y} ({k.fark} gün geç)")
            if k.miktar and k.teslim_miktar is not None:
                sapma = (k.teslim_miktar - k.miktar) / k.miktar
                if sapma < -tolerans:
                    uyar("Yüksek", "Eksik teslim", k, f"{binlik(k.teslim_miktar)} / {binlik(k.miktar)} {k.birim} teslim alındı (%{oran_yaz(-sapma)} eksik, "
                         f"tolerans %{oran_yaz(tolerans)}). Kalan miktar için tedarikçiden termin alın; kesim planını kontrol edin.")
                    k.durum = "Eksik teslim"
                elif sapma > tolerans:
                    uyar("Bilgi", "Fazla teslim", k, f"{binlik(k.teslim_miktar)} / {binlik(k.miktar)} {k.birim} (%{oran_yaz(sapma)} fazla)")
            continue
        # Tahmin
        imlec = k.gerceklesen[adlar[son_i]] if son_i >= 0 else (k.siparis or bugun)
        ilk = True
        for a, sure in k.asamalar[son_i + 1:]:
            onay_bekliyor = katla(a).startswith(BOYAMA) and k.labdip_sutunu and not k.labdip
            if katla(a).startswith(BOYAMA) and k.labdip:
                imlec = max(imlec, k.labdip)
            if onay_bekliyor:
                imlec = max(imlec, bugun)                  # onay gelmeden boyama başlamaz
            bitis = imlec + timedelta(days=sure)
            if ilk:
                k.mevcut_asama = a + (" (lab dip onayı bekleniyor)" if onay_bekliyor else "")
                if bitis < bugun:
                    uyar("Orta", "Aşama gecikti", k, f"{a}: {imlec:%d.%m.%Y} + {sure} gün = {bitis:%d.%m.%Y} bitmeliydi, {(bugun - bitis).days} gün geçti; "
                         "tedarikçiden güncel durum alın")
                    bitis = bugun
            if onay_bekliyor:
                uyar("Yüksek", "Lab dip onayı yok", k, f"{a} aşamasına gelinmiş / gelinecek ama lab dip onay tarihi boş. Onay gelmeden boyama "
                     f"başlamamalı; tahmin boyamanın bugün başlayacağını varsayar. Müşteriden onayı takip edin.")
            k.tahmini[a] = bitis
            imlec, ilk = bitis, False
        if son_i == len(adlar) - 1:
            k.mevcut_asama = "Yolda (sevk edildi)"
        k.tahmini_teslim = imlec
        if k.termin:
            k.fark = (k.tahmini_teslim - k.termin).days
            if bugun > k.termin:
                k.durum = "TERMİN GEÇTİ"
                uyar("Yüksek", "Termin geçti", k, f"Termin {k.termin:%d.%m.%Y}, {(bugun - k.termin).days} gün geçti; mevcut aşama {k.mevcut_asama}, "
                     f"tahmini teslim {k.tahmini_teslim:%d.%m.%Y}")
            elif k.fark > 0:
                k.durum = "Gecikecek"
                uyar("Orta", "Termin riski", k, f"Tahmini teslim {k.tahmini_teslim:%d.%m.%Y}, termin {k.termin:%d.%m.%Y} ({k.fark} gün geç)")
            else:
                k.durum = "Zamanında"
        else:
            k.durum = "Termin yok"

    # Sipariş etkisi
    etkiler = []
    gruplu = defaultdict(list)
    for k in kumaslar:
        gruplu[k.po].append(k)
    for po in sorted(set(gruplu) | set(siparisler)):
        sp = siparisler.get(po, {"no": po, "model": "", "musteri": "", "kesim": None, "sevk": None})
        lst = gruplu.get(po, [])
        if not lst:
            uyar("Orta", "Kumaş siparişi yok", None, f"{po}: konfeksiyon siparişine bağlı kumaş siparişi bulunamadı", po)
            continue
        hazirlar = [k for k in lst if k.hazir]
        kritik = max(hazirlar, key=lambda k: k.hazir) if hazirlar else None
        e = {"po": po, "model": sp["model"], "musteri": sp["musteri"], "kesim": sp["kesim"], "sevk": sp["sevk"], "kumas_sayisi": len(lst),
             "kritik": kritik, "hazir": kritik.hazir if kritik else None, "gereken": None, "kayma": None, "tahmini_sevk": None, "durum": "",
             "eksik": [k for k in lst if k.durum == "Eksik teslim"]}
        if kritik and sp["kesim"]:
            e["gereken"] = sp["kesim"] - timedelta(days=kontrol_gun)
            e["kayma"] = max((kritik.hazir - e["gereken"]).days, 0)
            bolluk = (e["gereken"] - kritik.hazir).days
            e["tahmini_sevk"] = sp["sevk"] + timedelta(days=e["kayma"]) if sp["sevk"] else None
            if e["kayma"] > 0:
                e["durum"] = "Kesim kayıyor"
                uyar("Kritik", "Kesim / sevk kayıyor", kritik, f"{po}: kritik kumaş {kritik.ad} tahmini {kritik.hazir:%d.%m.%Y} hazır; {kontrol_gun} gün giriş kontrolüyle "
                     f"kesim en erken {kritik.hazir + timedelta(days=kontrol_gun):%d.%m.%Y} (plan {sp['kesim']:%d.%m.%Y}) → {e['kayma']} gün kayma"
                     + (f"; kesim–sevk arası sabit kalırsa sevk {e['tahmini_sevk']:%d.%m.%Y} (plan {sp['sevk']:%d.%m.%Y})" if sp["sevk"] else ""), po)
            elif bolluk <= risk_gun:
                e["durum"] = "Riskli"
                uyar("Orta", "Kesim için az bolluk", kritik, f"{po}: kumaş {kritik.hazir:%d.%m.%Y} hazır, kesim için en geç {e['gereken']:%d.%m.%Y}; "
                     + ("bolluk yok" if bolluk == 0 else f"yalnız {bolluk} gün bolluk"), po)
            else:
                e["durum"] = "Zamanında"
        elif not sp["kesim"]:
            e["durum"] = "Kesim tarihi yok"
        if e["eksik"]:
            e["durum"] = (e["durum"] + " · " if e["durum"] else "") + "Eksik kumaş"
        etkiler.append(e)

    # Tedarikçi performansı
    perf = defaultdict(list)
    for k in kumaslar:
        perf[k.tedarikci or "—"].append(k)
    performans = []
    for ad, lst in sorted(perf.items()):
        teslim = [k for k in lst if k.teslim and k.termin]
        gec = [k.fark for k in teslim if k.fark and k.fark > 0]
        acik_gec = [k for k in lst if not k.teslim and k.durum in ("TERMİN GEÇTİ", "Gecikecek")]
        performans.append({"ad": ad, "siparis": len(lst), "teslim": len(teslim), "zamaninda": len(teslim) - len(gec),
                           "ort_gecikme": statistics.fmean(gec) if gec else 0, "acik_gec": len(acik_gec)})
    uyarilar.sort(key=lambda u: (SIRA[u["onem"]], u["po"], u["kumas"]))
    return {"uyarilar": uyarilar, "etkiler": etkiler, "performans": performans}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Kritik": "FDE2E1", "Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE", "TERMİN GEÇTİ": "FDE2E1", "Eksik teslim": "FDE2E1",
        "Gecikecek": "FFF4CE", "Zamanında": "E3F4E1", "Teslim alındı": "E3F4E1", "Kesim kayıyor": "FDE2E1", "Riskli": "FFF4CE"}
GERCEK = PatternFill("solid", fgColor="E3F4E1")
KONTROL = PatternFill("solid", fgColor="FFF4CE")
UST = Alignment(vertical="top", wrap_text=True)
TF = "DD.MM.YYYY"


def _baslik(ws, basliklar, genislik=()):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def _boya(c, anahtar):
    for parca in str(anahtar or "").split(" · "):
        if parca in RENK:
            c.fill = PatternFill("solid", fgColor=RENK[parca])
            return


def rapor_yaz(cikti: Path, kumaslar: list[Kumas], s: dict, bugun: date) -> None:
    wb = Workbook()
    se = wb.active
    se.title = "Sipariş Etkisi"
    _baslik(se, ["Sipariş", "Model", "Müşteri", "Kumaş Sayısı", "Kritik Kumaş", "Kumaş Hazır", "Kesim İçin En Geç", "Planlanan Kesim", "Kayma (gün)",
                 "Planlanan Sevk", "Tahmini Sevk", "Durum", "Aksiyon / Not"], (9, 8, 16, 8, 30, 11, 11, 11, 8, 11, 11, 20, 30))
    for e in s["etkiler"]:
        se.append([e["po"], e["model"], e["musteri"], e["kumas_sayisi"], e["kritik"].ad if e["kritik"] else "", e["hazir"], e["gereken"], e["kesim"],
                   e["kayma"], e["sevk"], e["tahmini_sevk"], e["durum"], ""])
        for j in (6, 7, 8, 10, 11):
            se.cell(se.max_row, j).number_format = TF
        _boya(se.cell(se.max_row, 12), e["durum"])
        se.cell(se.max_row, 13).fill = KONTROL
    se.append([])
    se.append([f"Durum tarihi {bugun:%d.%m.%Y}. Kayma, kesim ile sevk arasındaki sürenin değişmediği varsayımıyla sevke yansıtılır."])

    kd = wb.create_sheet("Kumaş Durumu")
    _baslik(kd, ["Kumaş Sipariş No", "Tedarikçi", "Kumaş", "Renk", "Tip", "Miktar", "Birim", "Bağlı Sipariş", "Sipariş Tarihi", "Termin",
                 "Mevcut Aşama", "Tahmini / Gerçek Teslim", "Fark (gün)", "Teslim Alınan", "Durum", "Tedarikçi Notu"],
            (11, 18, 18, 11, 7, 8, 5, 9, 11, 11, 18, 12, 8, 9, 14, 28))
    for k in kumaslar:
        kd.append([k.no, k.tedarikci, k.kumas, k.renk, k.tip, k.miktar, k.birim, k.po, k.siparis, k.termin, k.mevcut_asama, k.hazir, k.fark,
                   k.teslim_miktar, k.durum, ""])
        for j in (9, 10, 12):
            kd.cell(kd.max_row, j).number_format = TF
        if k.fark and k.fark > 0:
            kd.cell(kd.max_row, 13).font = Font(bold=True, color="C00000")
        _boya(kd.cell(kd.max_row, 15), k.durum)
        kd.cell(kd.max_row, 16).fill = KONTROL
    kd.auto_filter.ref = f"A1:P{kd.max_row}"

    am = wb.create_sheet("Aşama Matrisi")
    tum = list(dict.fromkeys(a for k in kumaslar for a, _ in k.asamalar))
    _baslik(am, ["Kumaş Sipariş No", "Kumaş", "Lab Dip Onayı"] + tum + ["Teslim"], (11, 26, 11) + (12,) * (len(tum) + 1))
    for k in kumaslar:
        am.append([k.no, f"{k.kumas} {k.renk}".strip(), k.labdip])
        am.cell(am.max_row, 3).number_format = TF
        for j, a in enumerate(tum, 4):
            c = am.cell(am.max_row, j)
            if a in k.gerceklesen:
                c.value, c.fill = k.gerceklesen[a], GERCEK
            elif a in k.tahmini:
                c.value, c.font = k.tahmini[a], Font(italic=True, color="7F7F7F")
            elif a not in [x for x, _ in k.asamalar]:
                c.value = "—"
            c.number_format = TF
        c = am.cell(am.max_row, len(tum) + 4, k.teslim or k.tahmini_teslim)
        c.number_format = TF
        if k.teslim:
            c.fill = GERCEK
        else:
            c.font = Font(italic=True, color="7F7F7F")
    am.append([])
    am.append(["Yeşil: gerçekleşen tarih · gri italik: tahmini bitiş (aşama standart süresiyle)"])

    tp = wb.create_sheet("Tedarikçi Performansı")
    _baslik(tp, ["Tedarikçi", "Kumaş Siparişi", "Teslim Alınan", "Zamanında", "Zamanında Oranı", "Geç Teslimde Ort. Gecikme (gün)", "Açık ve Gecikmede"],
            (22, 10, 10, 10, 10, 14, 12))
    for x in s["performans"]:
        tp.append([x["ad"], x["siparis"], x["teslim"], x["zamaninda"], x["zamaninda"] / x["teslim"] if x["teslim"] else None, round(x["ort_gecikme"], 1),
                   x["acik_gec"]])
        tp.cell(tp.max_row, 5).number_format = "0%"

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Sipariş", "Kumaş Sipariş No", "Tedarikçi", "Açıklama", "Aksiyon"], (9, 22, 9, 11, 18, 90, 24))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["po"], u["kumas"], u["tedarikci"], u["aciklama"], ""])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 6).alignment = UST
        uy.cell(uy.max_row, 7).fill = KONTROL
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(kumas_yolu: Path, sure_yolu: Path, cikti: Path, siparis_yolu: Path | None = None, bugun: date | None = None, kontrol_gun: int = 2,
             tolerans: float = 0.03) -> dict:
    bugun = bugun or date.today()
    sureler = sureleri_oku(sure_yolu)
    kumaslar = kumaslari_oku(kumas_yolu, sureler)
    siparisler = siparisleri_oku(siparis_yolu)
    s = analiz_et(kumaslar, siparisler, bugun, kontrol_gun, tolerans)
    rapor_yaz(cikti, kumaslar, s, bugun)
    return {**s, "kumaslar": kumaslar, "bugun": bugun}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Kumaş siparişlerinin örme / boyahane aşamalarını ve terminini takip eder, konfeksiyon sevkine etkisini hesaplar.")
    p.add_argument("--kumaslar", type=Path, default=ORNEK / "kumas_siparisleri.csv", help="Kumaş siparişleri ve aşama gerçekleşen tarihleri (.xlsx/.csv)")
    p.add_argument("--sureler", type=Path, default=ORNEK / "asama_sureleri.csv", help="Tip, Aşama, Süre (gün)")
    p.add_argument("--siparisler", type=Path, help="İsteğe bağlı: Sipariş No, Model, Planlanan Kesim Tarihi, Sevk Tarihi")
    p.add_argument("--bugun", help="Durum tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 08.10.2026)")
    p.add_argument("--kontrol-gun", type=int, default=2, help="Kumaş girişinden kesime kadar kontrol süresi (gün, varsayılan 2)")
    p.add_argument("--tolerans", type=float, default=3, help="Eksik / fazla teslim toleransı %% (varsayılan 3)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kumas_termin_takibi.xlsx")
    a = p.parse_args(argv)
    ornek = a.kumaslar == ORNEK / "kumas_siparisleri.csv"
    if ornek and a.siparisler is None:
        a.siparisler = ORNEK / "siparisler.csv"
    for y in (a.kumaslar, a.sureler, a.siparisler):
        if y is not None and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    bugun = tarih(a.bugun) if a.bugun else (date(2026, 10, 8) if ornek else None)
    if a.bugun and bugun is None:
        print("[X] --bugun GG.AA.YYYY biçiminde olmalı")
        return 1
    try:
        s = calistir(a.kumaslar, a.sureler, a.cikti, a.siparisler, bugun, a.kontrol_gun, a.tolerans / 100)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(s['kumaslar'])} kumaş siparişi · {len(s['etkiler'])} konfeksiyon siparişi · durum tarihi {s['bugun']:%d.%m.%Y}")
    for u in s["uyarilar"]:
        if u["onem"] in ("Kritik", "Yüksek"):
            print(f"[{'X' if u['onem'] == 'Kritik' else '!'}] {u['po']} {u['kumas']} {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
