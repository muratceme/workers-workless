import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
GUN = date(2026, 10, 10)


def f(no, gun, blok, kat, imalat, durum, karar=None, satir=0):
    return {"no": no, "tarih": date(2026, 10, gun), "blok": blok, "kat": kat, "imalat": imalat, "nokta": "", "durum": durum, "karar": date(2026, 10, karar) if karar else None,
            "kontrol": "", "aciklama": "", "satir": satir}


def veri():
    return [f("F1", 1, "A", "1", "Kalıp", "Onaylandı", 2), f("F2", 1, "A", "1", "Donatı", "Reddedildi", 2), f("F3", 3, "A", "1", "Donatı", "Onaylandı", 3), f("F4", 5, "A", "2", "Kalıp", "Bekliyor"),
            f("F5", 6, "B", "1", "Kalıp", "Reddedildi", 7), f("F6", 8, "B", "1", "Donatı", "Şartlı onay", 9)]


class TakipTesti(unittest.TestCase):
    def test_durum(self):
        r = main.takip(veri(), GUN)
        self.assertEqual(r["toplam"], {"talep": 6, "Onaylandı": 2, "Şartlı onay": 1, "Reddedildi": 2, "Bekliyor": 1, "red_oran": D("40.00"), "sure": D("0.80")})
        self.assertEqual(r["matris"], {("A", "1", "Donatı"): "Onaylandı (2. başvuru)", ("A", "1", "Kalıp"): "Onaylandı", ("A", "2", "Kalıp"): "Bekliyor", ("B", "1", "Donatı"): "Şartlı onay",
                                       ("B", "1", "Kalıp"): "Reddedildi"})
        self.assertEqual([(x["ad"], x["talep"], x["red_oran"]) for x in r["imalat"]], [("Donatı", 3, D("33.33")), ("Kalıp", 3, D("50.00"))])
        self.assertEqual((r["yerler"], [(b["no"], b["gun"]) for b in r["bekleyen"]]), ([("A", "1"), ("A", "2"), ("B", "1")], [("F4", 5)]))
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Bilgi", "Şartlı onay kapatılmadı", "B · 1 · Donatı"), ("Orta", "Red sonrası yeniden başvuru yok", "B · 1 · Kalıp"), ("Orta", "Uzun süredir bekleyen kontrol", "F4")])

    def test_yardimcilar(self):
        self.assertEqual([main._durum(x) for x in ("", "Onaylandı", "ŞARTLI ONAY", "Red", "uygun değil", "Uygun", "Açık", "belki")],
                         ["Bekliyor", "Onaylandı", "Şartlı onay", "Reddedildi", "Reddedildi", "Onaylandı", "Bekliyor", None])
        self.assertEqual(sorted(["10", "2", "Zemin", "B2", "1"], key=main._dogal), ["1", "2", "10", "B2", "Zemin"])
        self.assertEqual([u["tur"] for u in main.takip(veri(), GUN, 5)["uyarilar"]].count("Uzun süredir bekleyen kontrol"), 0)
        with self.assertRaises(ValueError):
            main.takip([], GUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "kontrol_formlari.csv", cikti, GUN)
            self.assertEqual(r["toplam"], {"talep": 40, "Onaylandı": 29, "Şartlı onay": 1, "Reddedildi": 8, "Bekliyor": 2, "red_oran": D("21.05"), "sure": D("0.97")})
            self.assertEqual([(u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Red sonrası yeniden başvuru yok", "A · 3 · Beton döküm öncesi"), ("Red sonrası yeniden başvuru yok", "B · Zemin · Tesisat (döşeme içi)"),
                              ("Uzun süredir bekleyen kontrol", "KF-101"), ("Şartlı onay kapatılmadı", "A · 1 · Beton döküm öncesi")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Durum Matrisi", "Özet", "Bekleyenler", "Tüm Formlar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "f.csv"
            y.write_text("ITP No;Tarih;Blok;Kat;İmalat;Sonuç;Karar Tarihi\n1;01.10.2026;A;1;Kalıp;Onay;02.10.2026\n2;x;A;1;Kalıp;;\n3;01.10.2026;A;1;Kalıp;belki;\n4;02.10.2026;A;1;Kalıp;Red;01.10.2026\n5;01.10.2026;;;;;\n",
                         encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["durum"], x["karar"]) for x in l], [u["tur"] for u in uy]),
                             ([("1", "Onaylandı", date(2026, 10, 2)), ("4", "Reddedildi", None)], ["Satır okunamadı", "Satır okunamadı", "Karar talepten önce"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--formlar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "10.10.2026", "--gun", "1"]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--gun", "-1"])], [2] * 2)
            self.assertEqual(main.main(["--formlar", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("ITP No;Blok\n1;A\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
