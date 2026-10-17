"""
İmalat Kontrol Formu Takibi — Gerisi Bende kod bloğu
Kalite Kontrol › Kalite Kontrol Mühendisi
İmalat kontrol taleplerini (ITP) ve onay durumlarını blok, kat ve imalat türü bazında takip eder.
  Güncel durum      = her blok · kat · imalat (· kontrol noktası) için en son talebin durumu
  Red oranı %       = reddedilen ÷ karar verilen (onay + şartlı onay + red) × 100
  Karar süresi      = karar (onay / red) tarihi − talep tarihi, gün
  Uyarı             = --gun (2) günden uzun süredir bekleyen talep · son durumu red olup yeniden başvurulmamış imalat ·
                      şartlı onayı kapatılmamış imalat (bilgi)
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"no": ("form no", "talep no", "itp no", "kontrol no", "no"), "tarih": ("talep tarihi", "tarih"), "blok": ("blok", "yapi", "bina"), "kat": ("kat", "seviye", "kot"),
         "imalat": ("imalat turu", "imalat", "is kalemi"), "nokta": ("kontrol noktasi", "kontrol asamasi", "asama"), "durum": ("durum", "sonuc", "onay durumu"),
         "karar": ("onay tarihi", "karar tarihi", "kontrol tarihi"), "kontrol": ("kontrol eden", "kontrol muhendisi", "onaylayan"), "aciklama": ("aciklama", "not", "red nedeni")}
DURUMLAR = ("Onaylandı", "Şartlı onay", "Reddedildi", "Bekliyor")


def _durum(x) -> str | None:
    k = katla(x)
    if not k or k.startswith("bekl") or k in ("acik", "talep edildi", "kontrol bekliyor"):
        return "Bekliyor"
    if "sartli" in k or "kosullu" in k:
        return "Şartlı onay"
    if k.startswith("red") or "uygun degil" in k or k.startswith("ret"):
        return "Reddedildi"
    if k.startswith("onay") or k in ("uygun", "kabul"):
        return "Onaylandı"
    return None


def _dogal(s: str) -> tuple:
    return tuple(int(p) if p.isdigit() else p for p in re.split(r"(\d+)", katla(s)))


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("tarih", "imalat")):
        imalat = metin(r["imalat"])
        if not imalat:
            continue
        t, d = tarih(r["tarih"]), _durum(r["durum"])
        k = tarih(r["karar"]) if metin(r["karar"]) else None
        kim = metin(r["no"]) or f"Satır {r['_satir']}"
        if t is None or d is None or (metin(r["karar"]) and k is None):
            uy.append(uyari("Yüksek", "Satır okunamadı", kim, f"Satır {r['_satir']}: talep tarihi, karar tarihi ya da durum ('{metin(r['durum'])}') okunamadı. Satır alınmadı."))
            continue
        if k and k < t:
            uy.append(uyari("Orta", "Karar talepten önce", kim, f"Talep {t:%d.%m.%Y}, karar {k:%d.%m.%Y}; süre hesabına girmedi."))
            k = None
        l.append({"no": metin(r["no"]), "tarih": t, "blok": metin(r["blok"]) or "-", "kat": metin(r["kat"]) or "-", "imalat": imalat, "nokta": metin(r["nokta"]), "durum": d, "karar": k,
                  "kontrol": metin(r["kontrol"]), "aciklama": metin(r["aciklama"]), "satir": r["_satir"]})
    return l, uy


def _ozet(l: list[dict]) -> dict:
    say = {d: sum(1 for x in l if x["durum"] == d) for d in DURUMLAR}
    karar = say["Onaylandı"] + say["Şartlı onay"] + say["Reddedildi"]
    sure = [(x["karar"] - x["tarih"]).days for x in l if x["karar"] and x["durum"] != "Bekliyor"]
    return {"talep": len(l), **{d: say[d] for d in DURUMLAR}, "red_oran": yuvarla(Decimal(say["Reddedildi"]) * 100 / karar) if karar else None,
            "sure": yuvarla(Decimal(sum(sure)) / len(sure)) if sure else None}


def takip(formlar: list[dict], gun: date, bekleme: int = 2) -> dict:
    if not formlar:
        raise ValueError("Kontrol talebi yok.")
    uy, son = [], {}
    for f in sorted(formlar, key=lambda f: (f["tarih"], f["karar"] or date.max, f["satir"])):
        anahtar = (katla(f["blok"]), katla(f["kat"]), katla(f["imalat"]), katla(f["nokta"]))
        onceki = son.get(anahtar)
        son[anahtar] = {**f, "deneme": (onceki["deneme"] + 1) if onceki else 1}
    for f in formlar:
        if f["durum"] == "Bekliyor" and (gun - f["tarih"]).days > bekleme:
            uy.append(uyari("Orta", "Uzun süredir bekleyen kontrol", f["no"] or f["imalat"], f"{f['blok']} · {f['kat']} · {f['imalat']}: {f['tarih']:%d.%m.%Y} talebi {(gun - f['tarih']).days} gündür bekliyor."))
    guncel = sorted(son.values(), key=lambda f: (_dogal(f["blok"]), _dogal(f["kat"]), katla(f["imalat"]), katla(f["nokta"])))
    for f in guncel:
        yer = f"{f['blok']} · {f['kat']} · {f['imalat']}" + (f" · {f['nokta']}" if f["nokta"] else "")
        if f["durum"] == "Reddedildi":
            uy.append(uyari("Orta", "Red sonrası yeniden başvuru yok", yer, f"{f['karar'] or f['tarih']:%d.%m.%Y} tarihinde reddedildi" + (f" ({f['aciklama']})" if f["aciklama"] else "") + "; düzeltme sonrası yeni talep yok."))
        elif f["durum"] == "Şartlı onay":
            uy.append(uyari("Bilgi", "Şartlı onay kapatılmadı", yer, "Şartlar yerine getirildikten sonra kapanış kontrolü kaydedilmemiş" + (f" ({f['aciklama']})." if f["aciklama"] else ".")))
    def kirilim(alan):
        g = {}
        for f in formlar:
            g.setdefault(katla(f[alan]), []).append(f)
        return sorted(({"ad": l[0][alan], **_ozet(l)} for l in g.values()), key=lambda x: _dogal(x["ad"]))
    imalatlar = sorted({f["imalat"] for f in guncel}, key=katla)
    yerler = sorted({(f["blok"], f["kat"]) for f in guncel}, key=lambda x: (_dogal(x[0]), _dogal(x[1])))
    matris = {}
    for f in guncel:
        hucre = f["durum"] + (f" ({f['deneme']}. başvuru)" if f["deneme"] > 1 else "")
        anahtar = (f["blok"], f["kat"], f["imalat"])
        matris[anahtar] = (matris[anahtar] + "; " if anahtar in matris else "") + (f"{f['nokta']}: " if f["nokta"] else "") + hucre
    bekleyen = sorted(({**f, "gun": (gun - f["tarih"]).days} for f in formlar if f["durum"] == "Bekliyor"), key=lambda f: (-f["gun"], f["no"]))
    return {"toplam": _ozet(formlar), "blok": kirilim("blok"), "kat": kirilim("kat"), "imalat": kirilim("imalat"), "guncel": guncel, "matris": matris, "imalatlar": imalatlar, "yerler": yerler,
            "bekleyen": bekleyen, "formlar": formlar, "uyarilar": uy}


DURUM_RENK = {"Onaylandı": YESIL, "Şartlı onay": SARI, "Reddedildi": KIRMIZI, "Bekliyor": MAVI}


def rapor_yaz(cikti: Path, r: dict, gun: date) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Durum Matrisi"
    _baslik(ws, ["Blok", "Kat"] + r["imalatlar"], (10, 10) + (20,) * len(r["imalatlar"]), "C2")
    for b, k in r["yerler"]:
        n = _satir(ws, [b, k] + [r["matris"].get((b, k, i)) for i in r["imalatlar"]], sar=range(3, len(r["imalatlar"]) + 3))
        for j, i in enumerate(r["imalatlar"]):
            v = r["matris"].get((b, k, i))
            if v:
                renk = next((DURUM_RENK[d] for d in ("Reddedildi", "Bekliyor", "Şartlı onay", "Onaylandı") if d in v), None)
                ws.cell(n, 3 + j).fill = PatternFill("solid", fgColor=renk)
    ws.append([])
    ws.append([f"Her hücre o blok · kat · imalatın en son talebinin durumudur. Rapor tarihi {gun:%d.%m.%Y}."])
    ws = wb.create_sheet("Özet")
    _baslik(ws, ["Kırılım", "Ad", "Talep", "Onaylandı", "Şartlı Onay", "Reddedildi", "Bekliyor", "Red Oranı %", "Ort. Karar Süresi (gün)"], (10, 20, 7, 9, 9, 10, 9, 10, 12))
    for ad, l in (("Blok", r["blok"]), ("Kat", r["kat"]), ("İmalat", r["imalat"]), ("TOPLAM", [{"ad": "", **r["toplam"]}])):
        for x in l:
            _satir(ws, [ad, x["ad"], x["talep"], x["Onaylandı"], x["Şartlı onay"], x["Reddedildi"], x["Bekliyor"], x["red_oran"], x["sure"]], para_s=(8, 9))
    ws = wb.create_sheet("Bekleyenler")
    _baslik(ws, ["Form No", "Talep Tarihi", "Bekleyen Gün", "Blok", "Kat", "İmalat", "Kontrol Noktası"], (12, 11, 9, 10, 10, 20, 22))
    for f in r["bekleyen"]:
        _satir(ws, [f["no"], f["tarih"], f["gun"], f["blok"], f["kat"], f["imalat"], f["nokta"]], tarih_s=(2,))
    ws = wb.create_sheet("Tüm Formlar")
    _baslik(ws, ["Form No", "Talep Tarihi", "Blok", "Kat", "İmalat", "Kontrol Noktası", "Durum", "Karar Tarihi", "Kontrol Eden", "Açıklama"], (12, 11, 10, 10, 20, 22, 12, 11, 16, 30), "B2")
    for f in sorted(r["formlar"], key=lambda f: (f["tarih"], f["no"])):
        _satir(ws, [f["no"], f["tarih"], f["blok"], f["kat"], f["imalat"], f["nokta"], f["durum"], f["karar"], f["kontrol"], f["aciklama"]], tarih_s=(2, 8), dolgu={7: DURUM_RENK[f["durum"]]}, sar=(10,))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, gun: date, bekleme: int = 2) -> dict:
    formlar, uy = oku(yol)
    r = takip(formlar, gun, bekleme)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, gun)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="İmalat kontrol taleplerini ve onay durumlarını blok, kat ve imalat türü bazında takip eder.")
    p.add_argument("--formlar", type=Path, help="Form No, Talep Tarihi, Blok, Kat, İmalat Türü, Kontrol Noktası, Durum, Onay Tarihi, Kontrol Eden, Açıklama")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 10.10.2026)")
    p.add_argument("--gun", type=int, default=2, help="Bu kadar günden uzun bekleyen talepler uyarılır (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "imalat_kontrol_takibi.xlsx")
    a = p.parse_args(argv)
    yol = a.formlar or ORNEK / "kontrol_formlari.csv"
    gun = tarih(a.tarih) if a.tarih else (date(2026, 10, 10) if a.formlar is None else date.today())
    if gun is None or a.gun < 0:
        print("[X] --tarih GG.AA.YYYY, --gun sıfır ya da üstü olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, gun, a.gun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {t['talep']} talep · onay {t['Onaylandı']} · şartlı {t['Şartlı onay']} · red {t['Reddedildi']} (%{t['red_oran']}) · bekleyen {t['Bekliyor']} · ortalama karar {t['sure']} gün")
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
