import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
GUN = date(2026, 9, 30)


def f(no, kurum, tur, t, tutar, kesinti, tahsil, tt=None):
    return {"no": no, "kurum": kurum, "tur": tur, "tarih": date(2026, *t), "tahsil_tarih": date(2026, *tt) if tt else None, "tutar": D(tutar), "kesinti": D(kesinti), "tahsil": D(tahsil)}


def veri():
    return [f("F1", "A", "SGK", (6, 1), 1000, 50, 950, (7, 31)), f("F2", "A", "SGK", (7, 1), 1000, 0, 1000, (8, 15)), f("F3", "A", "SGK", (9, 1), 2000, 100, 0),
            f("F4", "B", "Özel sigorta", (3, 1), 500, 100, 300, (6, 30)), f("F5", "b", "Özel sigorta", (9, 1), 500, 0, 0)]


def k(o):
    return (o["sayi"], o["fatura"], o["kesinti"], o["kesinti_oran"], o["net"], o["tahsil"], o["acik"], o["tahsil_oran"], o["sure"], {a: v for a, v in o["yas"].items() if v})


class RaporTesti(unittest.TestCase):
    def test_gostergeler(self):
        r = main.rapor(veri(), GUN)
        self.assertEqual([(o["kurum"],) + k(o) for o in r["kurumlar"]],
                         [("A", 3, D(4000), D(150), D("3.75"), D(3850), D(1950), D(1900), D("50.65"), D("52.31"), {"0-30": D(1900)}),
                          ("B", 2, D(1000), D(100), D("10.00"), D(900), D(300), D(600), D("33.33"), D("121.00"), {"0-30": D(500), "180+": D(100)})])
        self.assertEqual(k(r["toplam"]), (5, D(5000), D(250), D("5.00"), D(4750), D(2250), D(2500), D("47.37"), D("61.47"), {"0-30": D(2400), "180+": D(100)}))
        self.assertEqual([(o["tur"], o["fatura"]) for o in r["turler"]], [("Özel sigorta", D(1000)), ("SGK", D(4000))])
        self.assertEqual([(x["no"], x["gun"], x["acik"]) for x in r["acik"]], [("F4", 213, D(100)), ("F3", 29, D(1900)), ("F5", 29, D(500))])
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]), [("180 günden eski alacak", "B"), ("Kesinti oranı yüksek", "B"), ("Tahsil süresi uzun", "B")])

    def test_esik_ve_dilim(self):
        r = main.rapor(veri(), GUN, D(10), D("121"))
        self.assertEqual([u["tur"] for u in r["uyarilar"]], ["180 günden eski alacak"])
        yas = lambda gun: main.gosterge([f("X", "K", "", (1, 1), 100, 0, 0)], date(2026, 1, 1) + main.timedelta(days=gun))["yas"]   # noqa: E731
        self.assertEqual([next(a for a, v in yas(g).items() if v) for g in (0, 30, 31, 60, 61, 90, 91, 180, 181)], ["0-30", "0-30", "31-60", "31-60", "61-90", "61-90", "91-180", "91-180", "180+"])
        self.assertEqual([main._tur(x, y) for x, y in (("", "SGK İl Müdürlüğü"), ("ÖSS", "X"), ("", "Örnek Sigorta A.Ş."), ("Kurumsal", "Y"), ("", "Z"))], ["SGK", "Özel sigorta", "Özel sigorta", "Kurumsal", "Diğer"])
        with self.assertRaises(ValueError):
            main.rapor([], GUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "kurum_faturalari.csv", cikti, GUN)
            self.assertEqual([(o["kurum"], o["kesinti_oran"], o["tahsil_oran"], o["sure"]) for o in r["kurumlar"]],
                             [("SGK", D("2.76"), D("76.50"), D("68.10")), ("Sigorta Şirketi A", D("6.42"), D("88.91"), D("44.05")), ("Sigorta Şirketi B", D("3.18"), D("88.91"), D("37.05")),
                              ("Sigorta Şirketi C", D("8.54"), D("62.91"), D("112.75")), ("Kurumsal Anlaşma X", D("0.00"), D("79.61"), D("60.16"))])
            self.assertEqual((r["toplam"]["fatura"], r["toplam"]["kesinti_oran"], r["toplam"]["sure"]), (D("32575000.00"), D("3.70"), D("60.89")))
            self.assertEqual([(u["tur"], u["kim"]) for u in r["uyarilar"]], [("180 günden eski alacak", "Sigorta Şirketi C"), ("Kesinti oranı yüksek", "Sigorta Şirketi A"),
                                                                          ("Kesinti oranı yüksek", "Sigorta Şirketi C"), ("Tahsil süresi uzun", "Sigorta Şirketi C")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Kurumlar", "Kurum Türleri", "Açık Faturalar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "f.csv"
            y.write_text("İcmal No;Ödeyici;Tarih;Tutar;Red Tutarı;Ödenen;Ödeme Tarihi\n1;SGK;01.09.2026;1.000,00;10;990;20.09.2026\n1;SGK;01.09.2026;5;;;\n2;SGK;x;5;;;\n3;SGK;01.09.2026;5;9;;\n"
                         "4;SGK;01.09.2026;0;;;\n5;SGK;01.09.2026;100;;200;\n6;SGK;01.09.2026;100;;50;01.08.2026\n;SGK;01.09.2026;100;;;\n7;;01.09.2026;5;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["tur"], x["kesinti"], x["tahsil"], x["tahsil_tarih"]) for x in l], sorted(u["tur"] for u in uy)),
                             ([("1", "SGK", D(10), D(990), date(2026, 9, 20)), ("5", "SGK", D(0), D(200), None), ("6", "SGK", D(0), D(50), None), ("", "SGK", D(0), D(0), None)],
                              ["Fatura iki kez yazılmış", "Kesinti faturadan büyük", "Satır okunamadı", "Satır okunamadı", "Tahsil faturadan önce", "Tahsil tarihi yok", "Tahsilat net alacaktan fazla"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--faturalar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "30.09.2026", "--kesinti", "2", "--sure", "30"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--kesinti", "101"], ["--sure", "-1"])], [2] * 3)
            self.assertEqual(main.main(["--faturalar", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Ödeyici;Tutar\nA;5\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
