"""
Kurum Bazlı Gelir ve Kesinti Raporu — Workers / Workless kod bloğu
Anlaşmalı Kurumlar ve Faturalama › Anlaşmalı Kurumlar Müdürü
SGK ve özel sigorta şirketleri bazında fatura, tahsilat, kesinti oranı ve ortalama tahsil süresini raporlar.
  Kesinti oranı %        = kesinti tutarı ÷ fatura tutarı × 100
  Net alacak             = fatura tutarı − kesinti
  Açık bakiye            = net alacak − tahsil edilen
  Tahsil oranı %         = tahsil edilen ÷ net alacak × 100
  Ortalama tahsil süresi = Σ(tahsil edilen × (tahsil tarihi − fatura tarihi)) ÷ Σ tahsil edilen   (tutar ağırlıklı, gün)
  Yaşlandırma            = açık bakiye fatura tarihinden rapor tarihine geçen güne göre 0-30 / 31-60 / 61-90 / 91-180 / 180+
  Uyarı                  = kesinti oranı --kesinti (%5) üstünde · ortalama tahsil süresi --sure (90) gün üstünde · 180 günden eski açık bakiye
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"no": ("fatura no", "fatura numarasi", "fatura", "icmal no"), "kurum": ("kurum", "kurum adi", "sigorta sirketi", "anlasmali kurum", "odeyici"),
         "tur": ("kurum turu", "tur", "odeyici turu"), "tarih": ("fatura tarihi", "tarih", "donem sonu"), "tutar": ("fatura tutari", "tutar", "fatura toplami"),
         "kesinti": ("kesinti", "kesinti tutari", "red tutari", "kesilen tutar"), "tahsil": ("tahsil edilen", "tahsilat", "odenen", "tahsil tutari"),
         "tahsil_tarih": ("tahsil tarihi", "odeme tarihi", "tahsilat tarihi")}
DILIM = (("0-30", 30), ("31-60", 60), ("61-90", 90), ("91-180", 180), ("180+", None))
SIFIR = Decimal(0)


def _tur(x, kurum: str) -> str:
    k = katla(x) or katla(kurum)
    if "sgk" in k or "sosyal guvenlik" in k:
        return "SGK"
    if k in ("oss", "ozel sigorta", "ozel saglik sigortasi", "tss", "tamamlayici") or "sigorta" in k:
        return "Özel sigorta"
    return metin(x) or "Diğer"


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("kurum", "tarih", "tutar")):
        kurum, no = metin(r["kurum"]), metin(r["no"])
        if not kurum:
            continue
        t, tt = tarih(r["tarih"]), tarih(r["tahsil_tarih"]) if metin(r["tahsil_tarih"]) else None
        d = {a: para(r[a]) if metin(r[a]) else SIFIR for a in ("tutar", "kesinti", "tahsil")}
        kim = f"{kurum} · {no or r['_satir']}"
        if t is None or (metin(r["tahsil_tarih"]) and tt is None) or any(metin(r[a]) and (para(r[a]) is None or para(r[a]) < 0) for a in d) or d["tutar"] <= 0:
            uy.append(uyari("Yüksek", "Satır okunamadı", kim, f"Satır {r['_satir']}: tarih ya da tutarlardan biri okunamadı (tutarlar sıfır ya da üstü olmalı). Satır alınmadı."))
            continue
        if d["kesinti"] > d["tutar"]:
            uy.append(uyari("Yüksek", "Kesinti faturadan büyük", kim, f"Satır {r['_satir']}: kesinti {d['kesinti']:,.2f}, fatura {d['tutar']:,.2f}. Satır alınmadı."))
            continue
        if no and katla(no) in gorulen:
            uy.append(uyari("Orta", "Fatura iki kez yazılmış", kim, f"Satır {r['_satir']}: ikinci satır alınmadı."))
            continue
        if no:
            gorulen.add(katla(no))
        if d["tahsil"] > d["tutar"] - d["kesinti"]:
            uy.append(uyari("Orta", "Tahsilat net alacaktan fazla", kim, f"Tahsil edilen {d['tahsil']:,.2f}, net alacak {d['tutar'] - d['kesinti']:,.2f}."))
        if d["tahsil"] and tt is None:
            uy.append(uyari("Bilgi", "Tahsil tarihi yok", kim, "Tahsilat var ama tarihi yazılmamış; ortalama tahsil süresine girmedi."))
        if tt and tt < t:
            uy.append(uyari("Orta", "Tahsil faturadan önce", kim, f"Tahsil {tt:%d.%m.%Y}, fatura {t:%d.%m.%Y}; süre hesabına girmedi."))
            tt = None
        l.append({"no": no, "kurum": kurum, "tur": _tur(r["tur"], kurum), "tarih": t, "tahsil_tarih": tt, **d})
    return l, uy


def gosterge(l: list[dict], gun: date) -> dict:
    fatura, kesinti, tahsil = (sum((f[a] for f in l), SIFIR) for a in ("tutar", "kesinti", "tahsil"))
    net = fatura - kesinti
    sureli = [(f["tahsil"], (f["tahsil_tarih"] - f["tarih"]).days) for f in l if f["tahsil"] and f["tahsil_tarih"]]
    agirlik = sum((t for t, _ in sureli), SIFIR)
    yas = {ad: SIFIR for ad, _ in DILIM}
    for f in l:
        acik = f["tutar"] - f["kesinti"] - f["tahsil"]
        if acik > 0:
            g = (gun - f["tarih"]).days
            yas[next(ad for ad, ust in DILIM if ust is None or g <= ust)] += acik
    return {"sayi": len(l), "fatura": fatura, "kesinti": kesinti, "kesinti_oran": yuvarla(kesinti * 100 / fatura) if fatura else None, "net": net, "tahsil": tahsil,
            "acik": net - tahsil, "tahsil_oran": yuvarla(tahsil * 100 / net) if net else None, "sure": yuvarla(sum((t * s for t, s in sureli), SIFIR) / agirlik) if agirlik else None, "yas": yas}


def rapor(faturalar: list[dict], gun: date, kesinti_esik: Decimal = Decimal(5), sure_esik: Decimal = Decimal(90)) -> dict:
    if not faturalar:
        raise ValueError("Fatura kaydı yok.")
    uy, kurumlar, turler = [], {}, {}
    for f in faturalar:
        kurumlar.setdefault(katla(f["kurum"]), []).append(f)
        turler.setdefault(f["tur"], []).append(f)
    sonuc = []
    for k in sorted(kurumlar):
        l = kurumlar[k]
        o = {"kurum": l[0]["kurum"], "tur": l[0]["tur"], **gosterge(l, gun)}
        sonuc.append(o)
        if o["kesinti_oran"] is not None and o["kesinti_oran"] > kesinti_esik:
            uy.append(uyari("Orta", "Kesinti oranı yüksek", o["kurum"], f"{o['fatura']:,.2f} faturada {o['kesinti']:,.2f} kesinti (%{o['kesinti_oran']}; eşik %{kesinti_esik})."))
        if o["sure"] is not None and o["sure"] > sure_esik:
            uy.append(uyari("Orta", "Tahsil süresi uzun", o["kurum"], f"Ortalama tahsil süresi {o['sure']} gün (eşik {sure_esik})."))
        if o["yas"]["180+"] > 0:
            uy.append(uyari("Orta", "180 günden eski alacak", o["kurum"], f"Fatura tarihinden 180 günden fazla geçmiş açık bakiye {o['yas']['180+']:,.2f}."))
    sonuc.sort(key=lambda o: (-o["fatura"], katla(o["kurum"])))
    tur = [{"tur": t, **gosterge(l, gun)} for t, l in sorted(turler.items(), key=lambda x: katla(x[0]))]
    acik = sorted(({**f, "acik": f["tutar"] - f["kesinti"] - f["tahsil"], "gun": (gun - f["tarih"]).days} for f in faturalar if f["tutar"] - f["kesinti"] - f["tahsil"] > 0),
                  key=lambda f: (-f["gun"], katla(f["kurum"]), f["no"]))
    return {"kurumlar": sonuc, "turler": tur, "toplam": gosterge(faturalar, gun), "acik": acik, "uyarilar": uy}


BASLIK = ["Fatura Sayısı", "Fatura Tutarı", "Kesinti", "Kesinti %", "Net Alacak", "Tahsil Edilen", "Tahsil %", "Açık Bakiye", "Ort. Tahsil Süresi (gün)"] + [f"Açık {ad} gün" for ad, _ in DILIM]


def _deger(o: dict) -> list:
    return [o["sayi"], o["fatura"], o["kesinti"], o["kesinti_oran"], o["net"], o["tahsil"], o["tahsil_oran"], o["acik"], o["sure"]] + [o["yas"][ad] or None for ad, _ in DILIM]


def rapor_yaz(cikti: Path, r: dict, gun: date, kesinti_esik: Decimal, sure_esik: Decimal) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kurumlar"
    para_s = [3 + i for i in range(13)]
    _baslik(ws, ["Kurum", "Kurum Türü"] + BASLIK, (24, 13, 8, 15, 13, 8, 15, 15, 8, 15, 10, 12, 12, 12, 12, 12), "C2")
    for o in r["kurumlar"]:
        _satir(ws, [o["kurum"], o["tur"]] + _deger(o), para_s=para_s[1:],
               dolgu={6: KIRMIZI if o["kesinti_oran"] is not None and o["kesinti_oran"] > kesinti_esik else None, 11: KIRMIZI if o["sure"] is not None and o["sure"] > sure_esik else None,
                      16: KIRMIZI if o["yas"]["180+"] else None})
    k = _satir(ws, ["TOPLAM", ""] + _deger(r["toplam"]), para_s=para_s[1:])
    for h in ws[k]:
        h.font = Font(bold=True)
    ws.append([])
    ws.append([f"Rapor tarihi {gun:%d.%m.%Y}. Tahsil süresi tutar ağırlıklıdır. Yaşlandırma fatura tarihine göredir."])
    ws = wb.create_sheet("Kurum Türleri")
    _baslik(ws, ["Kurum Türü"] + BASLIK, (16, 8, 15, 13, 8, 15, 15, 8, 15, 10, 12, 12, 12, 12, 12))
    for o in r["turler"]:
        _satir(ws, [o["tur"]] + _deger(o), para_s=[x - 1 for x in para_s[1:]])
    ws = wb.create_sheet("Açık Faturalar")
    _baslik(ws, ["Kurum", "Fatura No", "Fatura Tarihi", "Geçen Gün", "Fatura Tutarı", "Kesinti", "Tahsil Edilen", "Açık Bakiye"], (24, 14, 11, 9, 14, 12, 14, 14))
    for f in r["acik"]:
        _satir(ws, [f["kurum"], f["no"], f["tarih"], f["gun"], f["tutar"], f["kesinti"], f["tahsil"], f["acik"]], para_s=(5, 6, 7, 8), tarih_s=(3,),
               dolgu={4: KIRMIZI if f["gun"] > 180 else SARI if f["gun"] > 90 else None})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, gun: date, kesinti_esik: Decimal = Decimal(5), sure_esik: Decimal = Decimal(90)) -> dict:
    faturalar, uy = oku(yol)
    r = rapor(faturalar, gun, kesinti_esik, sure_esik)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, gun, kesinti_esik, sure_esik)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="SGK ve özel sigorta şirketleri bazında fatura, tahsilat, kesinti oranı ve ortalama tahsil süresini raporlar.")
    p.add_argument("--faturalar", type=Path, help="Fatura No, Kurum, Kurum Türü, Fatura Tarihi, Fatura Tutarı, Kesinti, Tahsil Edilen, Tahsil Tarihi")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY; yaşlandırma için (varsayılan bugün; örnek veride 30.09.2026)")
    p.add_argument("--kesinti", default="5", help="Kesinti oranı bu yüzdenin üstündeyse uyarılır (varsayılan 5)")
    p.add_argument("--sure", default="90", help="Ortalama tahsil süresi bu kadar günü aşarsa uyarılır (varsayılan 90)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kurum_gelir_kesinti_raporu.xlsx")
    a = p.parse_args(argv)
    yol = a.faturalar or ORNEK / "kurum_faturalari.csv"
    gun = tarih(a.tarih) if a.tarih else (date(2026, 9, 30) if a.faturalar is None else date.today())
    kes, sure = para(a.kesinti), para(a.sure)
    if gun is None or kes is None or not 0 <= kes <= 100 or sure is None or sure < 0:
        print("[X] --tarih GG.AA.YYYY, --kesinti 0-100 arası, --sure sıfır ya da üstü olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, gun, kes, sure)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {len(r['kurumlar'])} kurum · {t['sayi']} fatura · {t['fatura']:,.2f} fatura · kesinti %{t['kesinti_oran']} · tahsil %{t['tahsil_oran']} · açık {t['acik']:,.2f}"
          + (f" · ortalama tahsil {t['sure']} gün" if t["sure"] is not None else ""))
    for o in r["turler"]:
        print(f"     {o['tur']}: fatura {o['fatura']:,.2f} · kesinti %{o['kesinti_oran']} · açık {o['acik']:,.2f}" + (f" · tahsil {o['sure']} gün" if o["sure"] is not None else ""))
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
