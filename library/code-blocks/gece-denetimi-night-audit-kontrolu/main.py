"""
Gece Denetimi (Night Audit) Kontrolü — Workers / Workless kod bloğu
Turizm ve Otelcilik › Ön Büro › Night Auditor (Gece Denetçisi)

Gün sonunda PMS'ten alınan raporları karşılaştırır:
  - Folyolar: çıkış yapmış ama bakiyesi kapanmamış (açık) folyo, çıkış günü geçmiş/gelmiş ama hâlâ konaklayan misafir,
    oda ücreti postalanmamış konaklayan, bakiye limiti aşımı.
  - Fiyat: postalanan oda ücreti ↔ fiyat kodu ve oda tipine göre fiyat listesi; ücretsiz (COMP / house use) odalar.
  - Oda durumu: PMS'te dolu ama kat hizmetleri boş gördü (skip) ya da tersi (sleep).
  - Kasa: kasiyer × ödeme tipi bazında sistem tahsilatı ↔ kasa sayımı / POS gün sonu.
  - Oda geliri: folyolardan hesaplanan oda geliri ↔ gelir raporundaki oda geliri (verilirse); doluluk, ADR, RevPAR.
İnternete bağlanmaz; PMS'e bağlanmaz.

Kullanım:
    python main.py                                         # örnek verilerle dener
    python main.py --folyolar folyolar.xlsx --fiyatlar fiyatlar.xlsx --oda-durumu hk.xlsx --tahsilat tahsilat.xlsx --kasa kasa.xlsx --tarih 07.10.2026
    python main.py ... --bakiye-limiti 15000 --kasa-tolerans 1 --rapor-oda-geliri 412500
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
SIFIR = Decimal(0)
KURUS = Decimal("0.01")


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def katla(s) -> str:
    s = kucuk(s)
    for a, b in zip("ıöüşğç", "iousgc"):
        s = s.replace(a, b)
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def para(x) -> Decimal:
    if x in (None, ""):
        return SIFIR
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("₺", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1 or re.fullmatch(r"-?\d{1,3}\.\d{3}", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return SIFIR


def yuvarla(x: Decimal) -> Decimal:
    return x.quantize(KURUS, rounding=ROUND_HALF_UP)


def tl(x: Decimal) -> str:
    return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d %H:%M:%S", "%d.%m.%Y %H:%M"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            continue
    return None


def tablo_oku(yol: Path, anahtarlar: tuple) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    satirlar = [r for r in satirlar if any(c not in (None, "") for c in r)]
    ak = {katla(a) for a in anahtarlar}
    bi = next((i for i, r in enumerate(satirlar[:10]) if any(katla(c) in ak for c in r)), 0)
    return satirlar[bi:]


def _bul(baslik, *adlar):
    b = [katla(x) for x in baslik]
    return next((b.index(katla(a)) for a in adlar if katla(a) in b), None)


def _al(r, i):
    return r[i] if i is not None and i < len(r) else None


def _str(r, i) -> str:
    return str(_al(r, i) or "").strip()


def oda_no(x) -> str:
    s = str(x or "").strip()
    return s[:-2] if s.endswith(".0") else s


def _oku(yol: Path, tanim: dict, zorunlu: tuple) -> tuple[list[list], dict]:
    s = tablo_oku(yol, tuple(a for k in zorunlu for a in tanim[k]))
    i = {k: _bul(s[0], *v) for k, v in tanim.items()}
    eksik = [k for k in zorunlu if i[k] is None]
    if eksik:
        raise SystemExit(f"{yol.name}: gerekli sütun bulunamadı ({', '.join(eksik)}). Başlıklar: {s[0]}")
    return [r for r in s[1:] if not katla(_al(r, i[zorunlu[0]])).startswith(("toplam", "genel"))], i


FOLYO = ("folyo no", "folio no", "folyo", "folio")
KASIYER = ("kasiyer", "cashier", "kullanıcı", "personel")
ODEME = ("ödeme tipi", "ödeme türü", "ödeme şekli", "payment type", "tahsilat tipi")


# ----------------------------------------------------------------------------
# Okuma
# ----------------------------------------------------------------------------

def durum_sinifi(s: str) -> str:
    k = katla(s)
    if any(w in k for w in ("no show", "gelmedi")):
        return "noshow"
    if any(w in k for w in ("iptal", "cancel")):
        return "iptal"
    if any(w in k for w in ("cikis", "checked out", "check out", "departed", "ayrildi")):
        return "cikis"
    if any(w in k for w in ("konak", "in house", "inhouse", "giris yapti", "checked in", "icerde")):
        return "konakliyor"
    if any(w in k for w in ("beklenen", "expected", "rezervasyon", "arrival")):
        return "beklenen"
    return k


def folyo_oku(yol: Path) -> list[dict]:
    satirlar, i = _oku(yol, {
        "folyo": FOLYO, "oda": ("oda no", "oda", "room", "room no"), "misafir": ("misafir", "misafir adı", "guest", "ad soyad"),
        "tip": ("oda tipi", "room type"), "kod": ("fiyat kodu", "rate code", "fiyat planı", "tarife"),
        "giris": ("giriş", "giriş tarihi", "arrival", "geliş"), "cikis": ("çıkış", "çıkış tarihi", "departure", "ayrılış"),
        "durum": ("durum", "status", "rezervasyon durumu"), "ucret": ("oda ücreti", "postalanan oda ücreti", "room charge", "günlük oda ücreti"),
        "bakiye": ("bakiye", "balance", "folyo bakiyesi"), "acente": ("acente", "kaynak", "agency")}, ("folyo", "durum"))
    return [{"folyo": _str(r, i["folyo"]), "oda": oda_no(_al(r, i["oda"])), "misafir": _str(r, i["misafir"]), "tip": _str(r, i["tip"]),
             "kod": _str(r, i["kod"]), "giris": tarih(_al(r, i["giris"])), "cikis": tarih(_al(r, i["cikis"])),
             "durum_ham": _str(r, i["durum"]), "durum": durum_sinifi(_str(r, i["durum"])),
             "ucret": para(_al(r, i["ucret"])) if i["ucret"] is not None and _al(r, i["ucret"]) not in (None, "") else None,
             "bakiye": para(_al(r, i["bakiye"])) if i["bakiye"] is not None else None, "acente": _str(r, i["acente"])}
            for r in satirlar if _al(r, i["folyo"])]


def fiyat_oku(yol: Path | None) -> dict[tuple[str, str], Decimal]:
    if not yol:
        return {}
    satirlar, i = _oku(yol, {"kod": ("fiyat kodu", "rate code", "fiyat planı", "tarife"), "tip": ("oda tipi", "room type"),
                             "fiyat": ("fiyat", "oda fiyatı", "günlük fiyat", "rate", "tutar")}, ("kod", "fiyat"))
    return {(katla(_al(r, i["kod"])), katla(_al(r, i["tip"]))): para(_al(r, i["fiyat"])) for r in satirlar if _al(r, i["kod"])}


def hk_oku(yol: Path | None) -> dict[str, str]:
    if not yol:
        return {}
    satirlar, i = _oku(yol, {"oda": ("oda no", "oda", "room", "room no"),
                             "durum": ("kat hizmetleri durumu", "hk durumu", "durum", "housekeeping status", "status")}, ("oda", "durum"))
    sonuc = {}
    for r in satirlar:
        k = katla(_al(r, i["durum"]))
        if any(w in k for w in ("ooo", "ariza", "out of order", "kullanim disi", "oos", "out of service")):
            d = "ooo"
        elif any(w in k for w in ("dolu", "occupied", "occ")):
            d = "dolu"
        elif any(w in k for w in ("bos", "vacant", "vac")):
            d = "bos"
        else:
            d = k
        sonuc[oda_no(_al(r, i["oda"]))] = d
    return sonuc


def tahsilat_oku(yol: Path | None) -> list[dict]:
    if not yol:
        return []
    satirlar, i = _oku(yol, {"folyo": FOLYO, "kasiyer": KASIYER, "tip": ODEME, "tutar": ("tutar", "tahsilat tutarı", "amount")},
                       ("tutar", "tip"))
    return [{"folyo": _str(r, i["folyo"]), "kasiyer": _str(r, i["kasiyer"]) or "-", "tip": _str(r, i["tip"]), "tutar": para(_al(r, i["tutar"]))}
            for r in satirlar if _al(r, i["tutar"]) not in (None, "")]


def kasa_oku(yol: Path | None) -> dict[tuple[str, str], Decimal]:
    if not yol:
        return {}
    satirlar, i = _oku(yol, {"kasiyer": KASIYER, "tip": ODEME,
                             "tutar": ("sayılan", "kasa sayımı", "sayım", "teslim edilen", "pos gün sonu", "tutar")}, ("tip", "tutar"))
    sonuc: dict = {}
    for r in satirlar:
        k = (_str(r, i["kasiyer"]) or "-", _str(r, i["tip"]))
        sonuc[k] = sonuc.get(k, SIFIR) + para(_al(r, i["tutar"]))
    return sonuc


# ----------------------------------------------------------------------------
# Kontrol
# ----------------------------------------------------------------------------

UCRETSIZ = ("comp", "house use", "hu", "ucretsiz", "kompliman", "komplimenter")


def kontrol(folyolar, fiyatlar, hk, tahsilat, kasa, gun: date, bakiye_limiti: Decimal, kasa_tol: Decimal, fiyat_tol: Decimal,
            rapor_oda_geliri: Decimal | None) -> tuple[list[dict], dict, list[dict]]:
    bulgular = []

    def ekle(seviye, alan, ref, kontrol_ad, aciklama, tutar=None):
        bulgular.append({"seviye": seviye, "alan": alan, "ref": ref, "kontrol": kontrol_ad, "aciklama": aciklama, "tutar": tutar})

    konaklayan = [f for f in folyolar if f["durum"] == "konakliyor"]
    oda_dolu: dict[str, list[dict]] = {}
    gorulen = set()
    for f in folyolar:
        ref = f"{f['oda'] or '-'} / {f['folyo']}"
        if f["folyo"] in gorulen:
            ekle("Hata", "Folyo", ref, "Mükerrer folyo", "Aynı folyo numarası birden fazla satırda")
        gorulen.add(f["folyo"])
        b = f["bakiye"] or SIFIR
        if f["durum"] == "cikis" and b != 0:
            ekle("Hata", "Folyo", ref, "Açık folyo", f"{f['misafir']} çıkış yapmış, bakiye {tl(b)} kapanmamış "
                 + ("(borç: tahsilat veya şirkete/acenteye aktarım gerekli)" if b > 0 else "(alacak: iade veya mahsup gerekli)"), b)
        if f["durum"] == "noshow" and b > 0:
            ekle("Dikkat", "Folyo", ref, "No-show bakiyesi", f"gelmeyen rezervasyonda bakiye {tl(b)}: no-show ücreti tahsil/aktarım kararı", b)
        if f["durum"] != "konakliyor":
            continue
        oda_dolu.setdefault(f["oda"], []).append(f)
        if f["cikis"] and f["cikis"] < gun:
            ekle("Hata", "Folyo", ref, "Çıkış tarihi geçmiş", f"{f['misafir']} çıkış {f['cikis']:%d.%m.%Y}, hâlâ konaklıyor görünüyor")
        elif f["cikis"] == gun:
            ekle("Dikkat", "Folyo", ref, "Bugün çıkışı var", f"{f['misafir']} bugün ayrılacaktı, çıkışı yapılmamış (uzatma mı, unutulan çıkış mı?)")
        if f["giris"] and f["giris"] > gun:
            ekle("Hata", "Folyo", ref, "Giriş tarihi ileri", f"giriş {f['giris']:%d.%m.%Y} > denetim günü")
        ucretsiz = katla(f["kod"]) in UCRETSIZ or any(katla(f["kod"]).startswith(u) for u in ("comp", "house"))
        if f["ucret"] is None or f["ucret"] == 0:
            if ucretsiz:
                ekle("Dikkat", "Fiyat", ref, "Ücretsiz oda", f"fiyat kodu {f['kod']}: yönetim onayı ve gerekçesi kontrol edilmeli")
            else:
                ekle("Hata", "Fiyat", ref, "Oda ücreti postalanmamış", f"{f['misafir']} konaklıyor ama bugünün oda ücreti yok (fiyat kodu {f['kod'] or '-'})")
        elif fiyatlar:
            beklenen = fiyatlar.get((katla(f["kod"]), katla(f["tip"]))) or fiyatlar.get((katla(f["kod"]), ""))
            if beklenen is None:
                ekle("Bilgi", "Fiyat", ref, "Fiyat kodu listede yok", f"fiyat kodu '{f['kod']}' / oda tipi '{f['tip']}' fiyat listesinde yok")
            elif abs(f["ucret"] - beklenen) > fiyat_tol:
                ekle("Yüksek", "Fiyat", ref, "Fiyat farkı", f"postalanan {tl(f['ucret'])}, {f['kod']} / {f['tip']} fiyatı {tl(beklenen)}",
                     f["ucret"] - beklenen)
        if bakiye_limiti and b > bakiye_limiti:
            ekle("Dikkat", "Folyo", ref, "Bakiye limiti", f"bakiye {tl(b)} > limit {tl(bakiye_limiti)}: ön provizyon / ara tahsilat", b)
    for oda, lst in oda_dolu.items():
        if oda and len(lst) > 1:
            ekle("Dikkat", "Oda", oda, "Aynı odada birden fazla folyo",
                 ", ".join(f["folyo"] for f in lst) + " — paylaşımlı oda (share) değilse çift satış")

    # Oda durumu
    if hk:
        for oda, d in sorted(hk.items()):
            pms = oda in oda_dolu
            if pms and d == "bos":
                ekle("Yüksek", "Oda", oda, "Skip", "PMS'te dolu, kat hizmetleri boş gördü: misafir ödemeden ayrılmış olabilir")
            elif not pms and d == "dolu":
                ekle("Yüksek", "Oda", oda, "Sleep", "PMS'te boş, kat hizmetleri dolu gördü: kayıtsız konaklama veya unutulan giriş")
            elif pms and d == "ooo":
                ekle("Hata", "Oda", oda, "Arızalı odada misafir", "Oda kullanım dışı (OOO) ama PMS'te konaklayan var")
        for oda in oda_dolu:
            if oda and oda not in hk:
                ekle("Bilgi", "Oda", oda, "Kat hizmetleri raporunda yok", "dolu oda, HK raporunda bulunamadı")

    # Tahsilat ↔ kasa
    folyo_set = {f["folyo"] for f in folyolar}
    for t in tahsilat:
        if t["folyo"] and t["folyo"] not in folyo_set:
            ekle("Dikkat", "Kasa", t["folyo"], "Folyosuz tahsilat", f"{t['kasiyer']} / {t['tip']} {tl(t['tutar'])}: folyo listesinde yok", t["tutar"])
    sistem: dict = {}
    for t in tahsilat:
        k = (t["kasiyer"], t["tip"])
        sistem[k] = sistem.get(k, SIFIR) + t["tutar"]
    kasa_sat = []
    if kasa:
        kasa_n = {(katla(a), katla(b)): (a, b, v) for (a, b), v in kasa.items()}
        sistem_n = {(katla(a), katla(b)): (a, b, v) for (a, b), v in sistem.items()}
        for k in sorted(set(kasa_n) | set(sistem_n)):
            a, b, sv = sistem_n.get(k, (None, None, SIFIR))
            a2, b2, kv = kasa_n.get(k, (None, None, SIFIR))
            fark = kv - sv
            kasa_sat.append({"kasiyer": a or a2, "tip": b or b2, "sistem": sv, "sayim": kv, "fark": fark})
            if abs(fark) > kasa_tol:
                ekle("Hata", "Kasa", f"{a or a2} / {b or b2}", "Kasa farkı",
                     f"sistem {tl(sv)}, sayım {tl(kv)}: {'fazla' if fark > 0 else 'eksik'} {tl(abs(fark))}", fark)

    # Gelir
    oda_geliri = sum((f["ucret"] or SIFIR for f in konaklayan), SIFIR)
    toplam_oda = len([d for d in hk.values() if d != "ooo"]) if hk else None
    dolu = len([o for o in oda_dolu if o])
    ucretli = len({f["oda"] for f in konaklayan if f["ucret"]})
    ozet = {"konaklayan": len(konaklayan), "dolu_oda": dolu, "toplam_oda": toplam_oda, "oda_geliri": oda_geliri,
            "adr": yuvarla(oda_geliri / ucretli) if ucretli else SIFIR,
            "doluluk": (Decimal(dolu) / toplam_oda * 100) if toplam_oda else None,
            "revpar": yuvarla(oda_geliri / toplam_oda) if toplam_oda else None,
            "acik_bakiye": sum((f["bakiye"] or SIFIR for f in folyolar if f["durum"] == "cikis" and f["bakiye"]), SIFIR),
            "tahsilat": sum((t["tutar"] for t in tahsilat), SIFIR), "kasa_farki": sum((k["fark"] for k in kasa_sat), SIFIR)}
    if rapor_oda_geliri is not None and abs(rapor_oda_geliri - oda_geliri) > Decimal("0.05"):
        ekle("Hata", "Gelir", "Oda geliri", "Gelir raporu farkı",
             f"folyolardan {tl(oda_geliri)}, gelir raporunda {tl(rapor_oda_geliri)}", rapor_oda_geliri - oda_geliri)
    return bulgular, ozet, kasa_sat


SEVIYE_SIRA = {"Hata": 0, "Yüksek": 1, "Dikkat": 2, "Bilgi": 3}


def calistir(folyo_yolu: Path, cikti: Path, gun: date, fiyat_yolu: Path | None = None, hk_yolu: Path | None = None,
             tahsilat_yolu: Path | None = None, kasa_yolu: Path | None = None, bakiye_limiti: float = 0, kasa_tolerans: float = 0,
             fiyat_tolerans: float = 0.5, rapor_oda_geliri: float | None = None) -> dict:
    folyolar = folyo_oku(folyo_yolu)
    bulgular, ozet, kasa_sat = kontrol(folyolar, fiyat_oku(fiyat_yolu), hk_oku(hk_yolu), tahsilat_oku(tahsilat_yolu), kasa_oku(kasa_yolu), gun,
                                       Decimal(str(bakiye_limiti)), Decimal(str(kasa_tolerans)), Decimal(str(fiyat_tolerans)),
                                       Decimal(str(rapor_oda_geliri)) if rapor_oda_geliri is not None else None)
    bulgular.sort(key=lambda b: (SEVIYE_SIRA[b["seviye"]], b["alan"], b["ref"]))
    _rapor(folyolar, bulgular, ozet, kasa_sat, gun, cikti)
    return {"bulgular": bulgular, "ozet": ozet, "kasa": kasa_sat, "folyolar": folyolar}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Hata": "F8C9C6", "Yüksek": "FDE2E1", "Dikkat": "FFF4CE", "Bilgi": "E8F0FE"}
PARA = "#,##0.00"


def _baslik(ws, satir=1):
    for c in ws[satir]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI


def _f(x):
    return None if x is None else float(x)


def _rapor(folyolar, bulgular, ozet, kasa_sat, gun, cikti):
    wb = Workbook()
    o = wb.active
    o.title = "Özet"
    o.append([f"Gece denetimi — {gun:%d.%m.%Y}"])
    o["A1"].font = Font(bold=True, size=12)
    satirlar = [("Konaklayan folyo", ozet["konaklayan"], None), ("Dolu oda", ozet["dolu_oda"], None),
                ("Satılabilir oda (OOO hariç)", ozet["toplam_oda"], None),
                ("Doluluk %", _f(ozet["doluluk"]), "0.0"), ("Oda geliri (folyolardan)", _f(ozet["oda_geliri"]), PARA),
                ("ADR (ücretli oda başına)", _f(ozet["adr"]), PARA), ("RevPAR", _f(ozet["revpar"]), PARA),
                ("Günün tahsilatı (sistem)", _f(ozet["tahsilat"]), PARA), ("Kasa farkı (sayım − sistem)", _f(ozet["kasa_farki"]), PARA),
                ("Çıkış yapmış açık folyo bakiyesi", _f(ozet["acik_bakiye"]), PARA)]
    for ad, v, fmt in satirlar:
        o.append([ad, v])
        if fmt:
            o.cell(o.max_row, 2).number_format = fmt
    o.append([])
    for s in SEVIYE_SIRA:
        o.append([s, sum(1 for b in bulgular if b["seviye"] == s)])
        o.cell(o.max_row, 1).fill = PatternFill("solid", fgColor=RENK[s])
    o.append([])
    o.append(["Not", "Gün kapatılmadan (date roll) önce Hata ve Yüksek bulguları çözün; çözülemeyenleri devir notuna yazın."])
    o.column_dimensions["A"].width = 34
    o.column_dimensions["B"].width = 18

    b = wb.create_sheet("Bulgular")
    b.append(["Seviye", "Alan", "Oda / Folyo / Kasa", "Kontrol", "Açıklama", "Tutar", "Yapılan İşlem"])
    _baslik(b)
    for x in bulgular:
        b.append([x["seviye"], x["alan"], x["ref"], x["kontrol"], x["aciklama"], _f(x["tutar"]), None])
        b.cell(b.max_row, 1).fill = PatternFill("solid", fgColor=RENK[x["seviye"]])
        b.cell(b.max_row, 6).number_format = PARA
    for j, w in enumerate((9, 8, 18, 26, 80, 13, 30), 1):
        b.column_dimensions[get_column_letter(j)].width = w
    b.freeze_panes = "D2"
    b.auto_filter.ref = b.dimensions

    k = wb.create_sheet("Kasa Mutabakatı")
    k.append(["Kasiyer", "Ödeme Tipi", "Sistem", "Sayım / POS", "Fark"])
    _baslik(k)
    for x in kasa_sat:
        k.append([x["kasiyer"], x["tip"], _f(x["sistem"]), _f(x["sayim"]), _f(x["fark"])])
        for c in (3, 4, 5):
            k.cell(k.max_row, c).number_format = PARA
        if x["fark"]:
            k.cell(k.max_row, 5).fill = PatternFill("solid", fgColor=RENK["Hata"])
    for j, w in enumerate((16, 18, 14, 14, 12), 1):
        k.column_dimensions[get_column_letter(j)].width = w

    f = wb.create_sheet("Folyolar")
    f.append(["Folyo No", "Oda", "Misafir", "Oda Tipi", "Fiyat Kodu", "Giriş", "Çıkış", "Durum", "Oda Ücreti", "Bakiye"])
    _baslik(f)
    for x in sorted(folyolar, key=lambda x: (x["oda"].zfill(6), x["folyo"])):
        f.append([x["folyo"], x["oda"], x["misafir"], x["tip"], x["kod"], x["giris"], x["cikis"], x["durum_ham"], _f(x["ucret"]), _f(x["bakiye"])])
        for c in (6, 7):
            f.cell(f.max_row, c).number_format = "DD.MM.YYYY"
        for c in (9, 10):
            f.cell(f.max_row, c).number_format = PARA
    for j, w in enumerate((11, 7, 22, 10, 11, 11, 11, 13, 12, 12), 1):
        f.column_dimensions[get_column_letter(j)].width = w
    f.freeze_panes = "C2"
    f.auto_filter.ref = f.dimensions
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ornek = BURASI / "ornek_veri"
    ap = argparse.ArgumentParser(description="Gün sonu folyo, fiyat, oda durumu ve kasa raporlarını karşılaştırır (night audit).")
    ap.add_argument("--folyolar", type=Path, default=ornek / "folyolar.csv",
                    help="Folyo No, Oda No, Misafir, Oda Tipi, Fiyat Kodu, Giriş, Çıkış, Durum, Oda Ücreti, Bakiye")
    ap.add_argument("--fiyatlar", type=Path, help="Fiyat Kodu, Oda Tipi, Fiyat")
    ap.add_argument("--oda-durumu", type=Path, help="Kat hizmetleri raporu: Oda No, Durum (Dolu / Boş / OOO)")
    ap.add_argument("--tahsilat", type=Path, help="Günün tahsilatları: Folyo No, Kasiyer, Ödeme Tipi, Tutar")
    ap.add_argument("--kasa", type=Path, help="Kasa sayımı / POS gün sonu: Kasiyer, Ödeme Tipi, Sayılan")
    ap.add_argument("--tarih", help="Denetim günü (gg.aa.yyyy); varsayılan bugün")
    ap.add_argument("--bakiye-limiti", type=float, default=0, help="Konaklayan misafir bakiyesi bu tutarı aşarsa uyar (0 = kapalı)")
    ap.add_argument("--kasa-tolerans", type=float, default=0, help="Kabul edilen kasa farkı, TL")
    ap.add_argument("--fiyat-tolerans", type=float, default=0.5, help="Kabul edilen fiyat farkı, TL (yuvarlama)")
    ap.add_argument("--rapor-oda-geliri", type=float, help="PMS gelir raporundaki günün oda geliri (karşılaştırma için)")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "gece_denetimi.xlsx")
    a = ap.parse_args(argv)
    if a.folyolar == ornek / "folyolar.csv":
        a.fiyatlar = a.fiyatlar or ornek / "fiyatlar.csv"
        a.oda_durumu = a.oda_durumu or ornek / "oda_durumu.csv"
        a.tahsilat = a.tahsilat or ornek / "tahsilat.csv"
        a.kasa = a.kasa or ornek / "kasa_sayim.csv"
        a.tarih = a.tarih or "07.10.2026"
        a.bakiye_limiti = a.bakiye_limiti or 15000
        if a.rapor_oda_geliri is None:
            a.rapor_oda_geliri = 25350
    gun = tarih(a.tarih) if a.tarih else date.today()
    if gun is None:
        raise SystemExit(f"Tarih anlaşılamadı: {a.tarih}")
    s = calistir(a.folyolar, a.cikti, gun, a.fiyatlar, a.oda_durumu, a.tahsilat, a.kasa, a.bakiye_limiti, a.kasa_tolerans,
                 a.fiyat_tolerans, a.rapor_oda_geliri)
    from collections import Counter
    say = Counter(b["seviye"] for b in s["bulgular"])
    print(f"[OK] {gun:%d.%m.%Y} · {s['ozet']['dolu_oda']} dolu oda · oda geliri {tl(s['ozet']['oda_geliri'])} TL · "
          + " · ".join(f"{k}: {say.get(k, 0)}" for k in SEVIYE_SIRA))
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
