import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def yk(adet, boy=120, en=80, yukseklik=100, kg=100, kat=1, ad="Y"):
    return {"ad": ad, "adet": adet, "boy": D(boy), "en": D(en), "yukseklik": D(yukseklik), "kg": D(kg), "kat": kat}


def ar(ad, uzunluk, kapasite, genislik="2.45", yukseklik="2.7"):
    return {"ad": ad, "uzunluk": D(uzunluk), "kapasite": D(kapasite), "genislik": D(genislik), "yukseklik": D(yukseklik)}


FILO = [ar("Tır", "13.6", 24000), ar("Kamyonet", "4.2", 3500, "2.1", "2.1"), ar("Kamyon", "7.2", 12000, "2.45", "2.5")]


class PlanTesti(unittest.TestCase):
    def test_ldm(self):
        y = main.planla([yk(10, yukseklik=150, kg=400)], FILO)["yukler"][0]
        self.assertEqual((y["zemin"], y["ldm"], y["hacim"], y["toplam_kg"]), (10, D("4.00"), D("14.40"), D(4000)))
        self.assertEqual([main.planla([yk(n, kat=2)], FILO)["yukler"][0]["ldm"] for n in (8, 7, 1)], [D("1.60"), D("1.60"), D("0.40")])           # 7 adet 2 katta 4 zemin yeri
        self.assertEqual(main.planla([yk(2, 300, 200, 180)], FILO)["yukler"][0]["ldm"], D("5.00"))
        self.assertEqual(main.planla([yk(33)], FILO)["toplam"]["ldm"], D("13.20"))                                                                 # 33 euro palet
        self.assertEqual(main.planla([yk(10)], FILO, D("2.45"))["toplam"]["ldm"], D("3.92"))
        self.assertEqual(main.planla([yk(1, kat=3, yukseklik=90)], FILO)["yukler"][0]["istif_yukseklik"], D(90))                                   # tek kap istiflenmez

    def test_arac_secimi(self):
        self.assertEqual([main.planla([yk(n, kg=k)], FILO)["secilen"]["ad"] for n, k in ((10, 100), (11, 100), (10, 400), (18, 100), (19, 100), (10, 1300))],
                         ["Kamyonet", "Kamyon", "Kamyon", "Kamyon", "Tır", "Tır"])                                                                 # 4,0 / 4,4 LDM; 4.000 kg; 7,2 / 7,6 LDM; 13.000 kg
        r = main.planla([yk(1, 100, 100, 220)], FILO)
        self.assertEqual((r["secilen"]["ad"], r["araclar"][0]["neden"]), ("Kamyon", ["Y: istifli yükseklik 220 cm araç iç yüksekliğini aşıyor"]))
        self.assertEqual(main.planla([yk(1, 300, 230)], FILO)["secilen"]["ad"], "Kamyon")                                                          # dar kenarı 230 cm: kamyonete sığmaz
        r = main.planla([yk(40)], FILO)                                                                                                            # 16 LDM
        self.assertEqual((r["secilen"], r["adet_arac"], r["buyuk"]["ad"]), (None, 2, "Tır"))
        self.assertEqual(main.planla([yk(20, kg=3000)], FILO)["adet_arac"], 3)                                                                     # 60.000 kg ÷ 24.000
        r = main.planla([yk(10, kg=400), yk(8, kg=250, kat=2), yk(2, 300, 200, 180, 1500)], FILO)
        self.assertEqual((r["toplam"], r["secilen"]["ad"], r["secilen"]["ldm_doluluk"], r["secilen"]["kg_doluluk"]),
                         ({"ldm": D("10.60"), "kg": D(9000), "hacim": D("38.88"), "adet": 20, "zemin": 16, "palet": D("26.50")}, "Tır", D("77.94"), D("37.50")))
        for y, a in (([], FILO), ([yk(1)], [])):
            with self.assertRaises(ValueError):
                main.planla(y, a)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "y.xlsx"
            r = main.calistir(ORNEK / "yukler.csv", ORNEK / "arac_tipleri.csv", cikti)
            self.assertEqual(([y["ldm"] for y in r["yukler"]], r["toplam"]["ldm"], r["toplam"]["kg"], r["toplam"]["hacim"], r["secilen"]["ad"].split(" (")[0], r["uyarilar"]),
                             ([D("4.00"), D("1.60"), D("5.00")], D("10.60"), D(9000), D("43.68"), "Tır - tenteli", []))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Yük Planı", "Araç Uygunluğu", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, a = Path(tmp) / "y.csv", Path(tmp) / "a.csv"
            y.write_text("Yük;Adet;Boy;En;Yükseklik;Ağırlık;İstiflenebilir\nA;2;120;80;100;1.250,5;Evet\nB;1,5;120;80;;;\nC;1;1,2;0,8;;10;üç\nD;4;120;80;50;10;3\n", encoding="utf-8")
            a.write_text("Araç Tipi;Yükleme Uzunluğu;Kapasite\nTır;13,6;24.000\nBozuk;uzun;1\n", encoding="utf-8")
            yl, al, uy = main.oku(y, a)
            self.assertEqual(([(x["ad"], x["adet"], x["kg"], x["kat"], x["yukseklik"]) for x in yl], [(x["ad"], x["uzunluk"], x["kapasite"], x["genislik"]) for x in al]),
                             ([("A", 2, D("1250.5"), 2, D(100)), ("C", 1, D(10), 1, D(0)), ("D", 4, D(10), 3, D(50))], [("Tır", D("13.6"), D(24000), None)]))
            self.assertEqual(sorted(u["tur"] for u in uy), ["Araç tipi okunamadı", "Yük satırı okunamadı", "Yükseklik yok", "Ölçü birimi şüpheli", "İstif katı okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--yukler", str(y), "--araclar", str(a), "--cikti", c]
            self.assertEqual(main.main(arg + ["--genislik", "2,45"]), 0)
            self.assertEqual(main.main(arg + ["--genislik", "0"]), 2)
            self.assertEqual(main.main(arg[:3] + [str(Path(tmp) / "yok.csv")]), 1)
            a.write_text("Araç Tipi;Yükleme Uzunluğu;Kapasite\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
