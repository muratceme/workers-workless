"""
Yük Planlama (LDM ve Palet) — Gerisi Bende kod bloğu
Karayolu Operasyon › Karayolu Operasyon Uzmanı

Yük ölçüleri, palet sayısı ve istiflenebilirliğe göre yükleme metresini (LDM), ağırlığı ve araç tipini hesaplar.
  Zemin yeri       = adet ÷ istif katı, yukarı yuvarlanır                 (istif katı 1: istiflenemez)
  LDM              = zemin yeri × boy (m) × en (m) ÷ araç genişliği       (--genislik, varsayılan 2,40 m)
  Euro palet eşdeğeri = LDM ÷ 0,4                                         (120 × 80 cm palet: 1,2 × 0,8 ÷ 2,4 = 0,4 LDM)
  Hacim (m³)       = adet × boy × en × yükseklik;   ağırlık = adet × birim ağırlık
  Araç seçimi      = araç tipleri dosyasında, toplam LDM ≤ yükleme uzunluğu, toplam ağırlık ≤ kapasite ve her kalem (en, boy, istifli yükseklik) sığan en kısa araç
  Sığmazsa         = en büyük araçtan kaç tane gerektiği: en büyük(LDM ÷ uzunluk, ağırlık ÷ kapasite), yukarı yuvarlanır
LDM bir alan yaklaşımıdır; gerçek yerleşim (kalemlerin yan yana sığması, dingil yükü, yük emniyeti) ayrıca kontrol edilmelidir. Araç ölçüleri sizin dosyanızdan gelir.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


YUK = {"ad": ("yuk", "kalem", "aciklama", "yuk tanimi", "referans"), "adet": ("adet", "kap adedi", "palet sayisi", "kap"), "boy": ("boy", "boy cm", "uzunluk"), "en": ("en", "en cm", "genislik"),
       "yuk": ("yukseklik", "yukseklik cm"), "kg": ("agirlik", "birim agirlik", "agirlik kg", "kg"), "kat": ("istif kati", "istiflenebilir", "istif", "kat")}
ARAC = {"ad": ("arac tipi", "arac", "tip"), "uzunluk": ("yukleme uzunlugu", "uzunluk", "ldm", "yukleme uzunlugu m"), "genislik": ("ic genislik", "genislik"), "yukseklik": ("ic yukseklik", "yukseklik"),
        "kapasite": ("kapasite", "kapasite kg", "tasima kapasitesi", "yuk kapasitesi")}
Z = Decimal(0)


def oku(yuk_yol: Path, arac_yol: Path) -> tuple[list[dict], list[dict], list[dict]]:
    uy, yukler, araclar = [], [], []
    for r in kayitlar(yuk_yol, YUK, ("adet", "boy", "en")):
        ad = metin(r["ad"]) or f"Satır {r['_satir']}"
        if not any(metin(r[a]) for a in ("adet", "boy", "en")):
            continue
        v = {a: para(r[a]) for a in ("adet", "boy", "en", "yuk", "kg")}
        if any(v[a] is None or v[a] <= 0 for a in ("adet", "boy", "en")) or v["adet"] != int(v["adet"]):
            uy.append(uyari("Yüksek", "Yük satırı okunamadı", ad, f"Adet '{metin(r['adet'])}', boy '{metin(r['boy'])}', en '{metin(r['en'])}'. Adet pozitif tam sayı, ölçüler cm olmalı; satır alınmadı."))
            continue
        k = katla(r["kat"])
        kat = 2 if k in ("e", "evet", "var") else 1 if k in ("", "h", "hayir", "yok") else para(r["kat"])
        if kat is None or kat < 1 or kat != int(kat):
            uy.append(uyari("Orta", "İstif katı okunamadı", ad, f"'{metin(r['kat'])}'. İstiflenemez (1 kat) sayıldı."))
            kat = 1
        for a, etiket in (("yuk", "Yükseklik"), ("kg", "Ağırlık")):
            if v[a] is None or v[a] < 0:
                uy.append(uyari("Orta", f"{etiket} yok", ad, f"'{metin(r[a])}'. 0 sayıldı; {'hacim ve yükseklik kontrolü' if a == 'yuk' else 'ağırlık kontrolü'} bu kalem için eksik."))
                v[a] = Z
        if max(v["boy"], v["en"]) < 10:
            uy.append(uyari("Orta", "Ölçü birimi şüpheli", ad, f"Boy {v['boy']:f}, en {v['en']:f}: ölçüler santimetre olmalı (metre yazılmış olabilir)."))
        yukler.append({"ad": ad, "adet": int(v["adet"]), "boy": v["boy"], "en": v["en"], "yukseklik": v["yuk"], "kg": v["kg"], "kat": int(kat)})
    for r in kayitlar(arac_yol, ARAC, ("ad", "uzunluk", "kapasite")):
        ad = metin(r["ad"])
        if not ad:
            continue
        u, k, g, y = (para(r[a]) for a in ("uzunluk", "kapasite", "genislik", "yukseklik"))
        if u is None or k is None or u <= 0 or k <= 0:
            uy.append(uyari("Yüksek", "Araç tipi okunamadı", ad, f"Yükleme uzunluğu '{metin(r['uzunluk'])}', kapasite '{metin(r['kapasite'])}'. Araç tipi kullanılmadı."))
            continue
        araclar.append({"ad": ad, "uzunluk": u, "kapasite": k, "genislik": g, "yukseklik": y})
    return yukler, araclar, uy


def planla(yukler: list[dict], araclar: list[dict], genislik: Decimal = Decimal("2.4")) -> dict:
    if not yukler:
        raise ValueError("Planlanacak yük yok.")
    if not araclar:
        raise ValueError("Araç tipi yok.")
    for y in yukler:
        y["zemin"] = math.ceil(Decimal(y["adet"]) / y["kat"])
        y["ldm"] = yuvarla(y["zemin"] * y["boy"] * y["en"] / 10000 / genislik)
        y["hacim"] = yuvarla(y["adet"] * y["boy"] * y["en"] * y["yukseklik"] / 1000000)
        y["toplam_kg"] = y["adet"] * y["kg"]
        y["istif_yukseklik"] = min(y["kat"], y["adet"]) * y["yukseklik"]
    ldm, kg = sum((y["ldm"] for y in yukler), Z), sum((y["toplam_kg"] for y in yukler), Z)
    sirali = sorted(araclar, key=lambda a: (a["uzunluk"], a["kapasite"]))
    for a in sirali:
        n = []
        if ldm > a["uzunluk"]:
            n.append(f"LDM {ldm:f} > {a['uzunluk']:f} m")
        if kg > a["kapasite"]:
            n.append(f"ağırlık {kg:f} > {a['kapasite']:f} kg")
        for y in yukler:
            if a["genislik"] and min(y["boy"], y["en"]) > a["genislik"] * 100:
                n.append(f"{y['ad']}: dar kenarı {min(y['boy'], y['en']):f} cm araç genişliğine sığmıyor")
            if max(y["boy"], y["en"]) > a["uzunluk"] * 100:
                n.append(f"{y['ad']}: uzun kenarı {max(y['boy'], y['en']):f} cm araç boyuna sığmıyor")
            if a["yukseklik"] and y["istif_yukseklik"] > a["yukseklik"] * 100:
                n.append(f"{y['ad']}: istifli yükseklik {y['istif_yukseklik']:f} cm araç iç yüksekliğini aşıyor")
        a["neden"], a["uygun"] = n, not n
        a["ldm_doluluk"], a["kg_doluluk"] = yuvarla(ldm / a["uzunluk"] * 100), yuvarla(kg / a["kapasite"] * 100)
    secilen = next((a for a in sirali if a["uygun"]), None)
    buyuk = sirali[-1]
    return {"yukler": yukler, "araclar": sirali, "secilen": secilen, "adet_arac": 1 if secilen else max(math.ceil(ldm / buyuk["uzunluk"]), math.ceil(kg / buyuk["kapasite"]), 1), "buyuk": buyuk,
            "toplam": {"ldm": ldm, "kg": kg, "hacim": sum((y["hacim"] for y in yukler), Z), "adet": sum(y["adet"] for y in yukler), "zemin": sum(y["zemin"] for y in yukler), "palet": yuvarla(ldm / Decimal("0.4"))},
            "genislik": genislik}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    t = r["toplam"]
    ws = wb.active
    ws.title = "Yük Planı"
    _baslik(ws, ["Yük", "Adet", "Boy (cm)", "En (cm)", "Yükseklik (cm)", "Birim Ağırlık (kg)", "İstif Katı", "Zemin Yeri", "LDM", "Toplam Ağırlık (kg)", "Hacim (m³)"], (34, 8, 9, 9, 10, 11, 8, 9, 10, 13, 11))
    for y in r["yukler"]:
        _satir(ws, [y["ad"], y["adet"], y["boy"], y["en"], y["yukseklik"], y["kg"], y["kat"], y["zemin"], y["ldm"], y["toplam_kg"], y["hacim"]], para_s=(9, 10, 11))
    _satir(ws, ["Toplam", t["adet"], "", "", "", "", "", t["zemin"], t["ldm"], t["kg"], t["hacim"]], para_s=(9, 10, 11))
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.append([])
    s = r["secilen"]
    ws.append([f"Önerilen araç: {s['ad']} (LDM doluluk %{s['ldm_doluluk']}, ağırlık doluluk %{s['kg_doluluk']})" if s else
               f"Tek araca sığmıyor: en büyük araç tipinden ({r['buyuk']['ad']}) yaklaşık {r['adet_arac']} araç gerekir."])
    ws.append([f"Euro palet eşdeğeri {t['palet']:f}. LDM = zemin yeri × boy × en ÷ {r['genislik']:f} m. LDM alan yaklaşımıdır; gerçek yerleşimi, dingil yükünü ve yük emniyetini ayrıca kontrol edin."])
    ws = wb.create_sheet("Araç Uygunluğu")
    _baslik(ws, ["Araç Tipi", "Yükleme Uzunluğu (m)", "Kapasite (kg)", "Uygun", "LDM Doluluk %", "Ağırlık Doluluk %", "Uygun Değilse Neden"], (26, 12, 12, 9, 11, 11, 80))
    for a in r["araclar"]:
        _satir(ws, [a["ad"], a["uzunluk"], a["kapasite"], "Evet" if a["uygun"] else "Hayır", a["ldm_doluluk"], a["kg_doluluk"], "; ".join(a["neden"])], para_s=(2, 5, 6),
               dolgu={4: YESIL if a is s else None if a["uygun"] else KIRMIZI}, sar=(7,))
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yuk_yol: Path, arac_yol: Path, cikti: Path, genislik: Decimal = Decimal("2.4")) -> dict:
    yukler, araclar, uy = oku(yuk_yol, arac_yol)
    r = planla(yukler, araclar, genislik)
    if not r["secilen"]:
        uy.append(uyari("Yüksek", "Tek araca sığmıyor", r["buyuk"]["ad"], f"Toplam {r['toplam']['ldm']:f} LDM, {r['toplam']['kg']:f} kg. Yaklaşık {r['adet_arac']} araç gerekir; yükü araçlara siz bölün (kod bölmez)."))
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Yükleme metresi (LDM), ağırlık ve uygun araç tipini hesaplar.")
    p.add_argument("--yukler", type=Path, default=ORNEK / "yukler.csv", help="Yük, Adet, Boy (cm), En (cm), Yükseklik (cm), Ağırlık (kg, birim), İstif Katı (1 = istiflenemez)")
    p.add_argument("--araclar", type=Path, default=ORNEK / "arac_tipleri.csv", help="Araç Tipi, Yükleme Uzunluğu (m), İç Genişlik (m), İç Yükseklik (m), Kapasite (kg)")
    p.add_argument("--genislik", default="2,4", help="LDM hesabında kullanılan araç genişliği, metre (varsayılan 2,40)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "yuk_plani.xlsx")
    a = p.parse_args(argv)
    g = para(a.genislik)
    if g is None or g <= 0:
        print("[X] --genislik pozitif sayı olmalı (metre)")
        return 2
    for y in (a.yukler, a.araclar):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.yukler, a.araclar, a.cikti, g)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t, s = r["toplam"], r["secilen"]
    print(f"[OK] {t['adet']} kap · {t['ldm']:f} LDM · {t['kg']:,.0f} kg · {t['hacim']:f} m³ · euro palet eşdeğeri {t['palet']:f}")
    print(f"[OK] Önerilen araç: {s['ad']} (LDM %{s['ldm_doluluk']}, ağırlık %{s['kg_doluluk']})" if s else f"[!] Tek araca sığmıyor: {r['buyuk']['ad']} tipinden yaklaşık {r['adet_arac']} araç")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
