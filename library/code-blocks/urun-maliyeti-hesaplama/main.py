"""
Ürün Maliyeti Hesaplama — Workers / Workless kod bloğu
Muhasebe › Maliyet Muhasebesi Uzmanı

Çok seviyeli reçete (ürün ağacı) ve rota üzerinden ürün birim maliyetini hesaplar:
  - Direkt ilk madde ve malzeme (DİMM): reçete miktarı × (1 + fire %) × birim maliyet; dövizli malzemeler verilen
    kurla TL'ye çevrilir.
  - Direkt işçilik (Dİ) ve genel üretim giderleri (GÜG): operasyon süresi (dk/adet) + hazırlık süresi / parti
    miktarı, iş merkezinin saatlik işçilik ve GÜG yükleme oranıyla; GÜG oranı yoksa işçiliğin yüzdesi olarak.
  - Yarı mamuller kendi maliyetleriyle (malzeme + işçilik + GÜG) üst ürüne aktarılır; maliyet ağacı, en pahalı
    bileşenler ve isteğe bağlı satış fiyatına göre brüt kâr marjı raporlanır.
İnternete bağlanmaz.

Kullanım:
    python main.py                                          # örnek verilerle dener
    python main.py --recete recete.xlsx --malzeme malzeme.xlsx --rota rota.xlsx --is-merkezi is_merkezi.xlsx
    python main.py ... --kur USD=41,20 EUR=48,10 --fiyat satis_fiyatlari.xlsx
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
SIFIR = Decimal(0)
ALTMIS = Decimal(60)


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def katla(s) -> str:
    s = kucuk(s)
    for a, b in zip("ıöüşğç", "iousgc"):
        s = s.replace(a, b)
    return re.sub(r"[^a-z0-9%]+", " ", s).strip()


def sayi(x, varsayilan: Decimal | None = SIFIR) -> Decimal | None:
    if x in (None, ""):
        return varsayilan
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("₺", "").replace("%", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return varsayilan


def tl(x, basamak: int = 2) -> str:
    return f"{x:,.{basamak}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


class Tablo:
    """Başlıklı tablo: sütunları Türkçe karakter/büyük harf farkını yok sayarak bulur."""

    def __init__(self, yol: Path):
        s = tablo_oku(yol)
        self.ad = yol.name
        self.bas = s[0] if s else []
        self.kb = [katla(x) for x in self.bas]
        self.satirlar = s[1:]

    def sutun(self, *adlar, zorunlu: str | None = None) -> int | None:
        i = next((self.kb.index(katla(a)) for a in adlar if katla(a) in self.kb), None)
        if i is None and zorunlu:
            raise SystemExit(f"{self.ad}: '{zorunlu}' sütunu bulunamadı. Başlıklar: {self.bas}")
        return i

    @staticmethod
    def al(r, i):
        return r[i] if i is not None and i < len(r) else None


# ----------------------------------------------------------------------------
# Veri okuma
# ----------------------------------------------------------------------------

def recete_oku(yol: Path) -> tuple[dict[str, list[dict]], dict[str, Decimal]]:
    """Dönüş: (ana ürün → [{kod, miktar, birim, fire}], ana ürün → reçete çıktı miktarı)."""
    t = Tablo(yol)
    i_ana = t.sutun("ana ürün", "ürün kodu", "mamul", "üst kod", "ana kod", zorunlu="Ana Ürün")
    i_bil = t.sutun("bileşen", "bileşen kodu", "malzeme kodu", "alt kod", zorunlu="Bileşen")
    i_mik = t.sutun("miktar", "bileşen miktarı", "kullanım miktarı", zorunlu="Miktar")
    i_bir = t.sutun("birim", "ölçü birimi")
    i_fire = t.sutun("fire %", "fire", "fire oranı", "fire yüzdesi")
    i_cikti = t.sutun("çıktı miktarı", "reçete miktarı", "taban miktar")
    recete: dict[str, list[dict]] = defaultdict(list)
    cikti: dict[str, Decimal] = {}
    for r in t.satirlar:
        ana, bil = str(t.al(r, i_ana) or "").strip(), str(t.al(r, i_bil) or "").strip()
        if not ana or not bil:
            continue
        recete[ana].append({"kod": bil, "miktar": sayi(t.al(r, i_mik)), "birim": str(t.al(r, i_bir) or "").strip(),
                            "fire": sayi(t.al(r, i_fire))})
        c = sayi(t.al(r, i_cikti), None)
        if c:
            cikti[ana] = c
    return dict(recete), cikti


def malzeme_oku(yol: Path, kurlar: dict[str, Decimal]) -> tuple[dict[str, dict], list[str]]:
    t = Tablo(yol)
    i_kod = t.sutun("malzeme kodu", "stok kodu", "kod", zorunlu="Malzeme Kodu")
    i_ad = t.sutun("malzeme adı", "stok adı", "ad", "açıklama")
    i_bir = t.sutun("birim", "ölçü birimi")
    i_mal = t.sutun("birim maliyet", "birim fiyat", "maliyet", "son alış fiyatı", "ortalama maliyet", zorunlu="Birim Maliyet")
    i_pb = t.sutun("para birimi", "döviz", "pb")
    malzeme, uyarilar = {}, []
    for r in t.satirlar:
        kod = str(t.al(r, i_kod) or "").strip()
        if not kod:
            continue
        pb = (str(t.al(r, i_pb) or "").strip().upper() or "TL").replace("TRY", "TL")
        tutar = sayi(t.al(r, i_mal))
        if pb != "TL":
            if pb not in kurlar:
                uyarilar.append(f"{kod}: {pb} kuru verilmedi (--kur {pb}=...); maliyet 0 alındı")
                tl_tutar = SIFIR
            else:
                tl_tutar = tutar * kurlar[pb]
        else:
            tl_tutar = tutar
        malzeme[kod] = {"ad": str(t.al(r, i_ad) or "").strip(), "birim": str(t.al(r, i_bir) or "").strip(),
                        "maliyet": tl_tutar, "orijinal": tutar, "pb": pb}
    return malzeme, uyarilar


def rota_oku(yol: Path | None) -> dict[str, list[dict]]:
    if not yol:
        return {}
    t = Tablo(yol)
    i_ur = t.sutun("ürün kodu", "ürün", "mamul", "ana ürün", zorunlu="Ürün Kodu")
    i_op = t.sutun("operasyon", "işlem", "operasyon adı")
    i_im = t.sutun("iş merkezi", "iş merkezi kodu", "makine", "istasyon", zorunlu="İş Merkezi")
    i_sure = t.sutun("süre dk", "birim süre dk", "süre (dk/adet)", "süre dk adet", "işlem süresi dk", "süre", zorunlu="Süre (dk/adet)")
    i_haz = t.sutun("hazırlık dk", "hazırlık süresi dk", "hazırlık süresi", "setup dk", "hazırlık")
    i_kisi = t.sutun("kişi sayısı", "operatör sayısı", "işçi sayısı")
    rota: dict[str, list[dict]] = defaultdict(list)
    for r in t.satirlar:
        ur = str(t.al(r, i_ur) or "").strip()
        if not ur:
            continue
        rota[ur].append({"op": str(t.al(r, i_op) or "").strip(), "im": str(t.al(r, i_im) or "").strip(),
                         "sure": sayi(t.al(r, i_sure)), "hazirlik": sayi(t.al(r, i_haz)), "kisi": sayi(t.al(r, i_kisi), Decimal(1))})
    return dict(rota)


def is_merkezi_oku(yol: Path | None) -> dict[str, dict]:
    if not yol:
        return {}
    t = Tablo(yol)
    i_im = t.sutun("iş merkezi", "iş merkezi kodu", "makine", "istasyon", zorunlu="İş Merkezi")
    i_isc = t.sutun("işçilik tl saat", "işçilik saat ücreti", "işçilik", "direkt işçilik tl saat", zorunlu="İşçilik (TL/saat)")
    i_gug = t.sutun("güg tl saat", "güg saat ücreti", "güg", "genel üretim gideri tl saat", "makine saat ücreti")
    return {str(t.al(r, i_im)).strip(): {"iscilik": sayi(t.al(r, i_isc)), "gug": sayi(t.al(r, i_gug), None)}
            for r in t.satirlar if t.al(r, i_im)}


def parti_oku(yol: Path | None) -> dict[str, Decimal]:
    if not yol:
        return {}
    t = Tablo(yol)
    i_ur = t.sutun("ürün kodu", "ürün", "kod", zorunlu="Ürün Kodu")
    i_p = t.sutun("parti miktarı", "parti", "üretim partisi", "lot", zorunlu="Parti Miktarı")
    return {str(t.al(r, i_ur)).strip(): sayi(t.al(r, i_p)) for r in t.satirlar if t.al(r, i_ur) and sayi(t.al(r, i_p))}


def fiyat_oku(yol: Path | None, kurlar: dict[str, Decimal]) -> dict[str, Decimal]:
    if not yol:
        return {}
    t = Tablo(yol)
    i_ur = t.sutun("ürün kodu", "ürün", "kod", zorunlu="Ürün Kodu")
    i_f = t.sutun("satış fiyatı", "net satış fiyatı", "fiyat", "liste fiyatı", zorunlu="Satış Fiyatı")
    i_pb = t.sutun("para birimi", "döviz", "pb")
    sonuc = {}
    for r in t.satirlar:
        ur = str(t.al(r, i_ur) or "").strip()
        if ur:
            pb = (str(t.al(r, i_pb) or "").strip().upper() or "TL").replace("TRY", "TL")
            sonuc[ur] = sayi(t.al(r, i_f)) * (kurlar.get(pb, Decimal(1)) if pb != "TL" else Decimal(1))
    return sonuc


# ----------------------------------------------------------------------------
# Hesap
# ----------------------------------------------------------------------------

def islem_maliyeti(ops: list[dict], im: dict[str, dict], parti: Decimal, gug_orani: Decimal | None,
                   uyarilar: list[str], urun: str) -> tuple[Decimal, Decimal, list[dict]]:
    """Ürünün kendi operasyonlarının birim işçilik ve GÜG'ü."""
    isc_t, gug_t, detay = SIFIR, SIFIR, []
    for o in ops:
        m = im.get(o["im"])
        if m is None:
            uyarilar.append(f"{urun}: '{o['im']}' iş merkezinin saat ücreti yok; işçilik/GÜG 0 alındı")
            m = {"iscilik": SIFIR, "gug": SIFIR}
        dk = o["sure"] + (o["hazirlik"] / parti if parti else SIFIR)
        saat = dk / ALTMIS
        isc = saat * m["iscilik"] * o["kisi"]
        if m["gug"] is not None:
            gug = saat * m["gug"]
        elif gug_orani is not None:
            gug = isc * gug_orani / 100
        else:
            gug = SIFIR
        isc_t += isc
        gug_t += gug
        detay.append({**o, "dk": dk, "iscilik": isc, "gug": gug})
    return isc_t, gug_t, detay


def calistir(recete_yolu: Path, malzeme_yolu: Path, cikti: Path, rota_yolu: Path | None = None,
             im_yolu: Path | None = None, parti_yolu: Path | None = None, fiyat_yolu: Path | None = None,
             kurlar: dict[str, Decimal] | None = None, gug_orani: float | None = None, varsayilan_parti: float = 1) -> dict:
    kurlar = kurlar or {}
    recete, cikti_mik = recete_oku(recete_yolu)
    malzeme, uyarilar = malzeme_oku(malzeme_yolu, kurlar)
    rota, im = rota_oku(rota_yolu), is_merkezi_oku(im_yolu)
    partiler, fiyatlar = parti_oku(parti_yolu), fiyat_oku(fiyat_yolu, kurlar)
    gug_o = Decimal(str(gug_orani)) if gug_orani is not None else None
    if rota and not im:
        uyarilar.append("Rota verildi ama iş merkezi ücretleri verilmedi; işçilik ve GÜG hesaplanamadı")

    sonuc: dict[str, dict] = {}
    yol: list[str] = []
    eksik: set[str] = set()

    def hesapla(kod: str) -> dict:
        if kod in sonuc:
            return sonuc[kod]
        if kod in yol:
            raise SystemExit("Reçetede döngü var: " + " → ".join(yol[yol.index(kod):] + [kod]))
        if kod not in recete:                               # satın alınan malzeme
            m = malzeme.get(kod)
            if m is None:
                eksik.add(kod)
                m = {"ad": "", "birim": "", "maliyet": SIFIR, "orijinal": SIFIR, "pb": "TL"}
            sonuc[kod] = {"kod": kod, "ad": m["ad"], "tur": "Malzeme", "birim": m["birim"], "malzeme": m["maliyet"],
                          "iscilik": SIFIR, "gug": SIFIR, "toplam": m["maliyet"], "satirlar": [], "ops": [], "pb": m["pb"],
                          "orijinal": m["orijinal"]}
            return sonuc[kod]
        yol.append(kod)
        taban = cikti_mik.get(kod, Decimal(1))             # reçete kaç adet için yazılmış
        mal = isc = gug = SIFIR
        satirlar = []
        for b in recete[kod]:
            alt = hesapla(b["kod"])
            brut = b["miktar"] * (1 + b["fire"] / 100) / taban
            satirlar.append({**b, "brut": brut, "birim_maliyet": alt["toplam"], "tutar": brut * alt["toplam"],
                             "mal": brut * alt["malzeme"], "isc": brut * alt["iscilik"], "gug": brut * alt["gug"], "tur": alt["tur"]})
            mal += brut * alt["malzeme"]
            isc += brut * alt["iscilik"]
            gug += brut * alt["gug"]
        parti = partiler.get(kod, Decimal(str(varsayilan_parti)))
        o_isc, o_gug, ops = islem_maliyeti(rota.get(kod, []), im, parti, gug_o, uyarilar, kod)
        yol.pop()
        ad = malzeme.get(kod, {}).get("ad", "")
        sonuc[kod] = {"kod": kod, "ad": ad, "tur": "Üretilen", "birim": malzeme.get(kod, {}).get("birim", ""),
                      "malzeme": mal, "iscilik": isc + o_isc, "gug": gug + o_gug, "toplam": mal + isc + o_isc + gug + o_gug,
                      "kendi_iscilik": o_isc, "kendi_gug": o_gug, "satirlar": satirlar, "ops": ops, "parti": parti}
        return sonuc[kod]

    alt_kodlar = {b["kod"] for bs in recete.values() for b in bs}
    mamuller = sorted(k for k in recete if k not in alt_kodlar)
    for k in sorted(recete):
        hesapla(k)
    if eksik:
        uyarilar.append(f"Maliyeti bulunamayan {len(eksik)} malzeme 0 alındı: {', '.join(sorted(eksik))}")
    for k in recete:
        if malzeme.get(k, {}).get("maliyet"):
            uyarilar.append(f"{k} hem reçetesi olan ürün hem malzeme listesinde maliyetli: reçeteden hesaplanan kullanıldı")
    for k in rota:
        if k not in recete:
            uyarilar.append(f"{k}: rotası var ama reçetesi yok; operasyon maliyeti hesaba katılmadı")
    for k, f in fiyatlar.items():
        if k in sonuc:
            sonuc[k]["fiyat"] = f
            sonuc[k]["marj"] = (f - sonuc[k]["toplam"]) / f * 100 if f else None
    _rapor(sonuc, mamuller, recete, uyarilar, kurlar, gug_o, cikti)
    return {"urunler": sonuc, "mamuller": mamuller, "uyarilar": uyarilar}


def patlat(sonuc: dict, kod: str, carpan: Decimal = Decimal(1), seviye: int = 0, cikis: list | None = None) -> list:
    """Maliyet ağacı satırları: (seviye, kod, birim başına brüt miktar, tutar)."""
    cikis = [] if cikis is None else cikis
    for s in sonuc[kod]["satirlar"]:
        m = s["brut"] * carpan
        cikis.append({"seviye": seviye + 1, "kod": s["kod"], "ad": sonuc[s["kod"]]["ad"], "tur": s["tur"], "birim": s["birim"],
                      "miktar": m, "fire": s["fire"], "birim_maliyet": s["birim_maliyet"], "tutar": m * s["birim_maliyet"]})
        if sonuc[s["kod"]]["tur"] == "Üretilen":
            patlat(sonuc, s["kod"], m, seviye + 1, cikis)
    return cikis


def malzeme_payi(sonuc: dict, kod: str) -> dict[str, Decimal]:
    """Mamulün birim maliyetindeki her satın alınan malzemenin toplam tutarı (tüm seviyeler)."""
    pay: dict[str, Decimal] = defaultdict(Decimal)
    for s in patlat(sonuc, kod):
        if s["tur"] == "Malzeme":
            pay[s["kod"]] += s["tutar"]
    return pay


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
PARA = "#,##0.00"
PARA4 = "#,##0.0000"


def _baslik(ws, satir=1):
    for c in ws[satir]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI


def _rapor(sonuc, mamuller, recete, uyarilar, kurlar, gug_o, cikti):
    wb = Workbook()
    o = wb.active
    o.title = "Birim Maliyetler"
    o.append(["Ürün Kodu", "Ürün Adı", "Tür", "DİMM (malzeme)", "Direkt İşçilik", "GÜG", "Birim Maliyet",
              "Malzeme %", "İşçilik %", "GÜG %", "Satış Fiyatı", "Brüt Kâr", "Brüt Kâr Marjı"])
    _baslik(o)
    sira = mamuller + sorted(k for k in recete if k not in mamuller)
    for k in sira:
        s = sonuc[k]
        t = s["toplam"] or Decimal(1)
        o.append([k, s["ad"], "Mamul" if k in mamuller else "Yarı mamul", float(s["malzeme"]), float(s["iscilik"]), float(s["gug"]),
                  float(s["toplam"]), float(s["malzeme"] / t), float(s["iscilik"] / t), float(s["gug"] / t),
                  float(s["fiyat"]) if "fiyat" in s else None,
                  float(s["fiyat"] - s["toplam"]) if "fiyat" in s else None,
                  float(s["marj"] / 100) if s.get("marj") is not None else None])
        n = o.max_row
        for c in (4, 5, 6, 7, 11, 12):
            o.cell(n, c).number_format = PARA
        for c in (8, 9, 10, 13):
            o.cell(n, c).number_format = "0.0%"
        if s.get("marj") is not None and s["marj"] < 0:
            o.cell(n, 13).fill = PatternFill("solid", fgColor="FDE2E1")
        o.cell(n, 7).font = Font(bold=True)
    for j, w in enumerate((14, 28, 11, 15, 15, 13, 15, 10, 10, 9, 13, 13, 13), 1):
        o.column_dimensions[get_column_letter(j)].width = w
    o.freeze_panes = "C2"
    if mamuller:
        g = BarChart()
        g.type, g.grouping, g.overlap = "bar", "stacked", 100
        g.title, g.height, g.width = "Mamul birim maliyet bileşenleri (TL)", 7, 18
        g.add_data(Reference(o, min_col=4, max_col=6, min_row=1, max_row=1 + len(mamuller)), titles_from_data=True)
        g.set_categories(Reference(o, min_col=1, min_row=2, max_row=1 + len(mamuller)))
        o.add_chart(g, f"B{len(sira) + 4}")

    a = wb.create_sheet("Maliyet Ağacı")
    a.append(["Mamul", "Seviye", "Bileşen", "Bileşen Adı", "Tür", "Birim", "Fire %", "Mamul Başına Brüt Miktar",
              "Bileşen Birim Maliyeti", "Mamul Başına Tutar"])
    _baslik(a)
    for m in mamuller:
        for s in patlat(sonuc, m):
            a.append([m, s["seviye"], ("  " * (s["seviye"] - 1)) + s["kod"], s["ad"], s["tur"], s["birim"], float(s["fire"]) or None,
                      float(s["miktar"]), float(s["birim_maliyet"]), float(s["tutar"])])
            n = a.max_row
            a.cell(n, 8).number_format = PARA4
            a.cell(n, 9).number_format = PARA
            a.cell(n, 10).number_format = PARA
            if s["tur"] == "Üretilen":
                for c in range(2, 11):
                    a.cell(n, c).font = Font(bold=True)
    for j, w in enumerate((12, 7, 18, 28, 10, 7, 7, 14, 14, 14), 1):
        a.column_dimensions[get_column_letter(j)].width = w
    a.freeze_panes = "B2"
    a.auto_filter.ref = a.dimensions

    op = wb.create_sheet("Operasyonlar")
    op.append(["Ürün Kodu", "Operasyon", "İş Merkezi", "Süre (dk/adet)", "Hazırlık (dk/parti)", "Parti", "Kişi",
               "Birim Süre (dk)", "İşçilik (TL/adet)", "GÜG (TL/adet)"])
    _baslik(op)
    for k in sira:
        for x in sonuc[k]["ops"]:
            op.append([k, x["op"], x["im"], float(x["sure"]), float(x["hazirlik"]), float(sonuc[k]["parti"]), float(x["kisi"]),
                       float(x["dk"]), float(x["iscilik"]), float(x["gug"])])
            for c in (8,):
                op.cell(op.max_row, c).number_format = "0.00"
            for c in (9, 10):
                op.cell(op.max_row, c).number_format = PARA
    for j, w in enumerate((14, 18, 12, 12, 14, 8, 6, 12, 14, 14), 1):
        op.column_dimensions[get_column_letter(j)].width = w

    p = wb.create_sheet("Malzeme Payları")
    p.append(["Mamul", "Malzeme", "Malzeme Adı", "Mamul Başına Tutar", "Mamul Maliyetindeki Pay", "Kümülatif Pay"])
    _baslik(p)
    for m in mamuller:
        pay = malzeme_payi(sonuc, m)
        kum = SIFIR
        t = sonuc[m]["toplam"] or Decimal(1)
        for kod, v in sorted(pay.items(), key=lambda i: -i[1]):
            kum += v
            p.append([m, kod, sonuc[kod]["ad"], float(v), float(v / t), float(kum / t)])
            p.cell(p.max_row, 4).number_format = PARA
            p.cell(p.max_row, 5).number_format = "0.0%"
            p.cell(p.max_row, 6).number_format = "0.0%"
    for j, w in enumerate((12, 14, 28, 16, 16, 12), 1):
        p.column_dimensions[get_column_letter(j)].width = w

    b = wb.create_sheet("Bilgi")
    satirlar = [["DİMM", "Direkt ilk madde ve malzeme: reçete miktarı × (1 + fire %) ÷ reçete çıktı miktarı × malzeme birim maliyeti"],
                ["Direkt işçilik", "(süre dk/adet + hazırlık dk / parti miktarı) ÷ 60 × işçilik TL/saat × kişi sayısı"],
                ["GÜG", "Birim süre (saat) × iş merkezi GÜG TL/saat" + (f"; GÜG oranı yoksa direkt işçiliğin %{gug_o:g}'i" if gug_o is not None else "")],
                ["Yarı mamul", "Kendi malzeme, işçilik ve GÜG'üyle hesaplanır; üst ürüne bu üç bileşen ayrı ayrı aktarılır"],
                ["Brüt kâr marjı", "(satış fiyatı − birim maliyet) ÷ satış fiyatı"],
                ["Not", "Standart (ön) maliyettir; gerçekleşen maliyetle farklar (fiyat, miktar, verimlilik) ayrıca analiz edilmelidir. "
                        "GÜG yükleme oranları dönemin normal kapasitesine göre belirlenmelidir."]]
    for pb, k in sorted(kurlar.items()):
        satirlar.append([f"Kur {pb}", tl(k, 4)])
    for u in uyarilar:
        satirlar.append(["Uyarı", u])
    for s in satirlar:
        b.append(s)
    b.column_dimensions["A"].width = 16
    b.column_dimensions["B"].width = 130
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def kur_coz(degerler: list[str]) -> dict[str, Decimal]:
    kurlar = {}
    for d in degerler or []:
        m = re.match(r"^([A-Za-z]{3})=([\d.,]+)$", d.strip())
        if not m:
            raise SystemExit(f"Kur biçimi PB=değer olmalı (ör. USD=41,20): {d}")
        kurlar[m.group(1).upper()] = sayi(m.group(2))
    return kurlar


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ornek = BURASI / "ornek_veri"
    ap = argparse.ArgumentParser(description="Reçete, malzeme maliyeti, rota ve iş merkezi ücretlerinden ürün birim maliyetini hesaplar.")
    ap.add_argument("--recete", type=Path, default=ornek / "recete.csv", help="Reçete: Ana Ürün, Bileşen, Miktar [, Birim, Fire %%, Çıktı Miktarı]")
    ap.add_argument("--malzeme", type=Path, default=ornek / "malzeme_maliyetleri.csv", help="Malzeme Kodu, Malzeme Adı, Birim, Birim Maliyet [, Para Birimi]")
    ap.add_argument("--rota", type=Path, help="Rota: Ürün Kodu, Operasyon, İş Merkezi, Süre (dk/adet) [, Hazırlık dk, Kişi Sayısı]")
    ap.add_argument("--is-merkezi", type=Path, help="İş Merkezi, İşçilik (TL/saat) [, GÜG (TL/saat)]")
    ap.add_argument("--parti", type=Path, help="Ürün Kodu, Parti Miktarı (hazırlık süresini dağıtmak için)")
    ap.add_argument("--fiyat", type=Path, help="Ürün Kodu, Satış Fiyatı [, Para Birimi]")
    ap.add_argument("--kur", nargs="*", help="Döviz kurları, ör. --kur USD=41,20 EUR=48,10")
    ap.add_argument("--gug-orani", type=float, help="İş merkezinde GÜG ücreti yoksa direkt işçiliğin yüzdesi olarak GÜG")
    ap.add_argument("--parti-varsayilan", type=float, default=1, help="Parti dosyasında olmayan ürünler için parti miktarı (varsayılan 1)")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "urun_maliyeti.xlsx")
    a = ap.parse_args(argv)
    if a.recete == ornek / "recete.csv" and a.malzeme == ornek / "malzeme_maliyetleri.csv":
        a.rota = a.rota or ornek / "rota.csv"
        a.is_merkezi = a.is_merkezi or ornek / "is_merkezleri.csv"
        a.parti = a.parti or ornek / "parti_miktarlari.csv"
        a.fiyat = a.fiyat or ornek / "satis_fiyatlari.csv"
        a.kur = a.kur or ["USD=41,20", "EUR=48,10"]
    s = calistir(a.recete, a.malzeme, a.cikti, a.rota, a.is_merkezi, a.parti, a.fiyat, kur_coz(a.kur), a.gug_orani, a.parti_varsayilan)
    for m in s["mamuller"]:
        u = s["urunler"][m]
        ek = f" · fiyat {tl(u['fiyat'])} · marj %{tl(u['marj'], 1)}" if u.get("marj") is not None else ""
        print(f"[OK] {m}: {tl(u['toplam'])} TL (malzeme {tl(u['malzeme'])}, işçilik {tl(u['iscilik'])}, GÜG {tl(u['gug'])}){ek}")
    for x in s["uyarilar"]:
        print(f"[!] {x}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
