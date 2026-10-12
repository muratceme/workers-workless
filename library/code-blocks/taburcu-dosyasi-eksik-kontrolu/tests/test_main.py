import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
RAPOR = date(2026, 10, 12)
KURAL = [{"belge": "Epikriz", "kosul": "", "adlar": {"epikriz", "cikis ozeti"}}, {"belge": "Onam", "kosul": "", "adlar": {"onam"}}, {"belge": "Ameliyat notu", "kosul": "Ameliyat", "adlar": {"ameliyat notu"}}]


def ds(protokol="P1", gun=-2, ameliyat=False, servis="S"):
    return {"protokol": protokol, "servis": servis, "hekim": "", "taburcu": None if gun is None else RAPOR + timedelta(days=gun), "kosul": {"ameliyat": ameliyat}}


def tek(belgeler, sure=None, **kw):
    return main.kontrol([ds(**kw)], {"p1": set(belgeler)}, KURAL, RAPOR, sure)["dosyalar"][0]


class KontrolTesti(unittest.TestCase):
    def test_gereken_ve_eksik(self):
        d = tek({"epikriz", "onam"})
        self.assertEqual((d["gereken"], d["eksik"], d["durum"], d["gecen"]), (["Epikriz", "Onam"], [], "Tam", 2))
        d = tek({"epikriz", "onam"}, ameliyat=True)
        self.assertEqual((d["gereken"], d["eksik"], d["durum"]), (["Epikriz", "Onam", "Ameliyat notu"], ["Ameliyat notu"], "Eksik"))
        self.assertEqual(tek({"cikis ozeti", "onam", "ameliyat notu"}, ameliyat=True)["durum"], "Tam")           # diğer ad da sayılır
        self.assertEqual(tek(set())["eksik"], ["Epikriz", "Onam"])
        self.assertEqual(main.kontrol([ds()], {}, KURAL, RAPOR)["dosyalar"][0]["eksik"], ["Epikriz", "Onam"])     # hiç belgesi olmayan dosya

    def test_sure(self):
        self.assertEqual([tek({"onam"}, 3, gun=g)["durum"] for g in (-3, -4)], ["Eksik", "Süre aşıldı"])           # tam 3. gün henüz aşılmadı
        self.assertEqual((tek({"onam"}, gun=-40)["durum"], tek({"onam"}, 3, gun=None)["durum"], tek({"epikriz", "onam"}, 3, gun=-40)["durum"]), ("Eksik", "Eksik", "Tam"))

    def test_ozetler(self):
        l = [ds("A", -1, True, "Cerrahi"), ds("B", -9, True, "Cerrahi"), ds("C", -5, servis="Dahiliye"), ds("D", -2, servis="Dahiliye")]
        b = {"a": {"epikriz", "onam", "ameliyat notu"}, "b": {"onam"}, "c": {"epikriz"}, "d": {"epikriz", "onam", "diyet"}, "z": {"epikriz"}}
        r = main.kontrol(l, b, KURAL, RAPOR, 3)
        self.assertEqual([(d["protokol"], d["durum"], d["eksik"]) for d in r["dosyalar"]], [("B", "Süre aşıldı", ["Epikriz", "Ameliyat notu"]), ("C", "Süre aşıldı", ["Onam"]), ("D", "Tam", []), ("A", "Tam", [])])   # aynı durumda taburcusu eski olan önce
        self.assertEqual([(s["servis"], s["dosya"], s["tam"], s["asildi"], s["tamlik"], s["eksik_belge"], s["en_cok"]) for s in r["servis"]],
                         [("Cerrahi", 2, 1, 1, D("50.00"), 2, "Epikriz"), ("Dahiliye", 2, 1, 1, D("50.00"), 1, "Onam")])
        self.assertEqual([(x["belge"], x["kosul"], x["gereken"], x["eksik"], x["oran"]) for x in r["belge"]], [("Epikriz", "Hepsi", 4, 1, D("25.00")), ("Onam", "Hepsi", 4, 1, D("25.00")), ("Ameliyat notu", "Ameliyat", 2, 1, D("50.00"))])
        self.assertEqual((r["say"], r["tamlik"], r["listede_yok"], r["taninmayan"]), ({"dosya": 4, "tam": 2, "eksik": 2, "asildi": 2, "eksik_belge": 3}, D("50.00"), ["z"], ["diyet"]))
        with self.assertRaises(ValueError):
            main.kontrol([], {}, KURAL, RAPOR)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "taburcu_listesi.csv", ORNEK / "belge_listesi.csv", ORNEK / "belge_kurallari.csv", cikti, RAPOR, 3)
            self.assertEqual([(d["protokol"], d["durum"], d["gecen"], d["eksik"]) for d in r["dosyalar"]],
                             [("P-1005", "Süre aşıldı", 10, ["Epikriz", "Aydınlatılmış onam formu"]), ("P-1002", "Süre aşıldı", 4, ["Aydınlatılmış onam formu", "Konsültasyon formu"]),
                              ("P-1003", "Eksik", 3, ["Epikriz"]), ("P-1001", "Tam", 7, []), ("P-1004", "Tam", 2, [])])
            self.assertEqual((r["say"], r["tamlik"], [u["tur"] for u in r["uyarilar"]]),
                             ({"dosya": 5, "tam": 2, "eksik": 3, "asildi": 2, "eksik_belge": 5}, D("40.00"), ["Belgesi var, taburcu listesinde yok", "Kuralda olmayan belge türleri"]))
            self.assertEqual([(x["belge"].split()[0], x["gereken"], x["eksik"]) for x in r["belge"]], [("Epikriz", 5, 2), ("Aydınlatılmış", 5, 2), ("Ameliyat", 2, 0), ("Konsültasyon", 2, 1)])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Dosya Durumu", "Servis Özeti", "Belge Özeti", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            t, b, k = Path(tmp) / "t.csv", Path(tmp) / "b.csv", Path(tmp) / "k.csv"
            k.write_text("Belge;Koşul;Diğer Adlar\nEpikriz;;Çıkış özeti / Özet\nAmeliyat notu;Ameliyat;\n;Hepsi;\n", encoding="utf-8")
            t.write_text("Dosya No;Klinik;Taburcu Tarihi;AMELİYAT\nA;X;10.10.2026;evet\nA;X;10.10.2026;hayır\nB;;dün;\n", encoding="utf-8")
            b.write_text("Protokol No;Belge Türü\nA;ÖZET\nA;\n", encoding="utf-8")
            kr = main.kural_oku(k)
            self.assertEqual([(x["belge"], x["kosul"], x["adlar"]) for x in kr], [("Epikriz", "", {"epikriz", "cikis ozeti", "ozet"}), ("Ameliyat notu", "Ameliyat", {"ameliyat notu"})])
            d, bl, uy = main.oku(t, b, kr)
            self.assertEqual(([(x["protokol"], x["servis"], x["taburcu"], x["kosul"]) for x in d], bl),
                             ([("A", "X", date(2026, 10, 10), {"ameliyat": True}), ("B", "Belirtilmemiş", None, {"ameliyat": False})], {"a": {"ozet"}}))
            self.assertEqual(sorted(u["tur"] for u in uy), ["Mükerrer protokol", "Taburcu tarihi okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--taburcu", str(t), "--belgeler", str(b), "--kurallar", str(k), "--cikti", c, "--tarih", "12.10.2026"]
            self.assertEqual(main.main(arg + ["--sure-gun", "2"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(arg + ["--sure-gun", "birkaç"]), 2)
            self.assertEqual(main.main(arg[:-1] + ["dün"]), 2)
            k.write_text("Belge;Koşul\nEpikriz;\nDiyaliz formu;Diyaliz\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)                                         # koşul sütunu taburcu listesinde yok
            k.write_text("Belge;Koşul\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--taburcu", str(Path(tmp) / "yok.csv")]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
