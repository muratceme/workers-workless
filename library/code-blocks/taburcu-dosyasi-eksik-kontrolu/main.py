"""
Taburcu Dosyası Eksik Kontrolü — Gerisi Bende kod bloğu
Tıbbi Sekreterlik › Yatan Hasta Tıbbi Sekreteri

Taburcu olan hastaların dosyalarında onam, epikriz, ameliyat notu ve konsültasyon belgelerinin eksiksizliğini kontrol eder.
  Gereken belgeler = kural dosyasındaki her belge; "Koşul" boş veya "Hepsi" ise her dosyada, bir sütun adıysa (ör. Ameliyat) o sütunu "Evet" olan dosyalarda aranır
  Belge var        = belge listesinde o protokol numarasıyla, kuraldaki adla veya "Diğer Adlar"dan biriyle kayıt bulunması (büyük-küçük harf ve Türkçe karakter ayrımı yapılmaz)
  Dosya durumu     = eksik yoksa Tam;  eksik varsa Eksik;  eksik var ve taburcudan beri geçen gün > --sure-gun ise Süre aşıldı (süre verilmezse bu durum kullanılmaz)
  Tamlık %         = tam dosya ÷ dosya sayısı × 100
Hangi belgenin zorunlu olduğunu kurumunuz belirler; örnek kural dosyası bir başlangıçtır, mevzuat listesi değildir. Kod belgenin varlığına bakar, içeriğine ve imzasına bakmaz.
Hasta adı ve kimlik numarası gerekmez; protokol numarası yeterlidir.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"protokol": ("protokol no", "protokol", "dosya no", "yatis no", "basvuru no"), "servis": ("servis", "klinik", "bolum"), "hekim": ("hekim", "doktor", "sorumlu hekim"),
         "taburcu": ("taburcu tarihi", "cikis tarihi", "taburcu")}
BELGE = {"protokol": SUTUN["protokol"], "belge": ("belge", "belge turu", "belge adi", "evrak", "dokuman")}
KURAL = {"belge": BELGE["belge"], "kosul": ("kosul", "sart", "ne zaman"), "diger": ("diger adlar", "diger adlari", "es anlamli", "takma adlar")}
EVET = {"e", "evet", "x", "1", "var", "true", "yapildi"}
HEPSI = {"", "hepsi", "tum", "tumu", "her dosya", "her zaman"}


def kural_oku(yol: Path) -> list[dict]:
    kurallar = []
    for r in kayitlar(yol, KURAL, ("belge",)):
        b = metin(r["belge"])
        if b:
            kosul = metin(r["kosul"])
            kurallar.append({"belge": b, "kosul": "" if katla(kosul) in HEPSI else kosul, "adlar": {katla(b)} | {katla(x) for x in re.split(r"[,|/]", metin(r["diger"])) if katla(x)}})
    if not kurallar:
        raise ValueError(f"{yol.name}: kural dosyasında belge yok.")
    return kurallar


def oku(taburcu_yol: Path, belge_yol: Path, kurallar: list[dict]) -> tuple[list[dict], dict, list[dict]]:
    uy, dosyalar, belgeler = [], {}, defaultdict(set)
    kosullar = sorted({k["kosul"] for k in kurallar if k["kosul"]})
    satirlar = kayitlar(taburcu_yol, SUTUN | {f"k_{katla(k)}": (katla(k),) for k in kosullar}, ("protokol",))
    baslik = next((r for r in tablo_oku(taburcu_yol)[:15] if any(katla(c) in SUTUN["protokol"] for c in r)), [])
    eksik = [k for k in kosullar if katla(k) not in {katla(c) for c in baslik}]
    if eksik:
        raise ValueError(f"{taburcu_yol.name}: kural dosyasındaki koşul sütunu yok: {', '.join(eksik)}. Taburcu listesine bu adla Evet / Hayır sütunu ekleyin.")
    for r in satirlar:
        p = metin(r["protokol"])
        if not p:
            continue
        kim = f"{p} (satır {r['_satir']})"
        if katla(p) in dosyalar:
            uy.append(uyari("Orta", "Mükerrer protokol", kim, "Taburcu listesinde aynı protokol numarası birden fazla; yalnız ilki alındı."))
            continue
        t = tarih(r["taburcu"])
        if t is None and metin(r["taburcu"]):
            uy.append(uyari("Orta", "Taburcu tarihi okunamadı", kim, f"'{metin(r['taburcu'])}'. Geçen gün hesaplanmadı."))
        dosyalar[katla(p)] = {"protokol": p, "servis": metin(r["servis"]) or "Belirtilmemiş", "hekim": metin(r["hekim"]), "taburcu": t,
                              "kosul": {katla(k): katla(r[f"k_{katla(k)}"]) in EVET for k in kosullar}}
    for r in kayitlar(belge_yol, BELGE, ("protokol", "belge")):
        p, b = metin(r["protokol"]), metin(r["belge"])
        if p and b:
            belgeler[katla(p)].add(katla(b))
    return list(dosyalar.values()), dict(belgeler), uy


def kontrol(dosyalar: list[dict], belgeler: dict, kurallar: list[dict], rapor: date, sure_gun: int | None = None) -> dict:
    if not dosyalar:
        raise ValueError("Kontrol edilecek taburcu dosyası yok.")
    belge_ozet = {k["belge"]: {"belge": k["belge"], "kosul": k["kosul"] or "Hepsi", "gereken": 0, "eksik": 0} for k in kurallar}
    for d in dosyalar:
        var = belgeler.get(katla(d["protokol"]), set())
        d["gereken"], d["eksik"] = [], []
        for k in kurallar:
            if k["kosul"] and not d["kosul"].get(katla(k["kosul"])):
                continue
            d["gereken"].append(k["belge"])
            belge_ozet[k["belge"]]["gereken"] += 1
            if not var & k["adlar"]:
                d["eksik"].append(k["belge"])
                belge_ozet[k["belge"]]["eksik"] += 1
        d["gecen"] = (rapor - d["taburcu"]).days if d["taburcu"] else None
        d["durum"] = "Tam" if not d["eksik"] else "Süre aşıldı" if sure_gun is not None and d["gecen"] is not None and d["gecen"] > sure_gun else "Eksik"
    servis = {}
    for d in dosyalar:
        s = servis.setdefault(katla(d["servis"]), {"servis": d["servis"], "dosya": 0, "tam": 0, "eksik_belge": 0, "asildi": 0, "sayac": defaultdict(int)})
        s["dosya"] += 1
        s["tam"] += d["durum"] == "Tam"
        s["asildi"] += d["durum"] == "Süre aşıldı"
        s["eksik_belge"] += len(d["eksik"])
        for b in d["eksik"]:
            s["sayac"][b] += 1
    for s in servis.values():
        s["tamlik"] = yuvarla(Decimal(s["tam"]) / s["dosya"] * 100)
        s["en_cok"] = max(s["sayac"], key=lambda b: (s["sayac"][b], b)) if s["sayac"] else ""
    for b in belge_ozet.values():
        b["oran"] = yuvarla(Decimal(b["eksik"]) / b["gereken"] * 100) if b["gereken"] else None
    taninan = set().union(*(k["adlar"] for k in kurallar))
    protokoller = {katla(d["protokol"]) for d in dosyalar}
    sira = ("Süre aşıldı", "Eksik", "Tam")
    tam = sum(d["durum"] == "Tam" for d in dosyalar)
    return {"dosyalar": sorted(dosyalar, key=lambda d: (sira.index(d["durum"]), -(d["gecen"] or 0), d["protokol"])), "servis": sorted(servis.values(), key=lambda s: (s["tamlik"], s["servis"])),
            "belge": list(belge_ozet.values()), "rapor": rapor, "sure": sure_gun, "say": {"dosya": len(dosyalar), "tam": tam, "eksik": len(dosyalar) - tam,
                                                                                       "asildi": sum(d["durum"] == "Süre aşıldı" for d in dosyalar), "eksik_belge": sum(len(d["eksik"]) for d in dosyalar)},
            "tamlik": yuvarla(Decimal(tam) / len(dosyalar) * 100), "listede_yok": sorted(p for p in belgeler if p not in protokoller),
            "taninmayan": sorted({b for bs in belgeler.values() for b in bs} - taninan)}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    renk = {"Süre aşıldı": KIRMIZI, "Eksik": SARI, "Tam": YESIL}
    ws = wb.active
    ws.title = "Dosya Durumu"
    _baslik(ws, ["Durum", "Protokol No", "Servis", "Hekim", "Taburcu Tarihi", "Geçen Gün", "Gereken Belge", "Eksik Belge Sayısı", "Eksik Belgeler", "Tamamlandı (siz doldurun)"], (13, 16, 22, 22, 12, 9, 9, 9, 60, 18))
    for d in r["dosyalar"]:
        _satir(ws, [d["durum"], d["protokol"], d["servis"], d["hekim"], d["taburcu"], d["gecen"], len(d["gereken"]), len(d["eksik"]) or None, ", ".join(d["eksik"]), ""], tarih_s=(5,),
               dolgu={1: renk[d["durum"]]}, sar=(9,))
    ws.append([])
    ws.append([f"Rapor tarihi {r['rapor']:%d.%m.%Y}. " + (f"Dosya tamamlama süresi {r['sure']} gün. " if r["sure"] is not None else "Tamamlama süresi verilmedi; 'Süre aşıldı' durumu kullanılmadı. ")
               + "Kontrol belgenin kayıtlı olmasına bakar; içeriği, imzası ve tarihi ayrıca incelenmelidir."])
    ws = wb.create_sheet("Servis Özeti")
    _baslik(ws, ["Servis", "Dosya", "Tam", "Eksik", "Süresi Aşılan", "Tamlık %", "Eksik Belge Sayısı", "En Çok Eksik Belge"], (24, 8, 8, 8, 10, 10, 10, 36))
    for s in r["servis"]:
        _satir(ws, [s["servis"], s["dosya"], s["tam"], s["dosya"] - s["tam"], s["asildi"] or None, s["tamlik"], s["eksik_belge"], s["en_cok"]], para_s=(6,), dolgu={5: KIRMIZI if s["asildi"] else None})
    y = r["say"]
    _satir(ws, ["Toplam", y["dosya"], y["tam"], y["eksik"], y["asildi"] or None, r["tamlik"], y["eksik_belge"], ""], para_s=(6,))
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws = wb.create_sheet("Belge Özeti")
    _baslik(ws, ["Belge", "Koşul", "Gereken Dosya", "Eksik", "Eksik %"], (40, 18, 10, 8, 9))
    for b in r["belge"]:
        _satir(ws, [b["belge"], b["kosul"], b["gereken"], b["eksik"], b["oran"]], para_s=(5,), dolgu={4: SARI if b["eksik"] else None})
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(taburcu_yol: Path, belge_yol: Path, kural_yol: Path, cikti: Path, rapor: date, sure_gun: int | None = None) -> dict:
    kurallar = kural_oku(kural_yol)
    dosyalar, belgeler, uy = oku(taburcu_yol, belge_yol, kurallar)
    r = kontrol(dosyalar, belgeler, kurallar, rapor, sure_gun)
    for p in r["listede_yok"]:
        uy.append(uyari("Orta", "Belgesi var, taburcu listesinde yok", p, "Belge listesinde bu protokol için kayıt var ama taburcu listesinde yok. Protokol numarası yazımını veya dönem aralığını kontrol edin."))
    if r["taninmayan"]:
        uy.append(uyari("Bilgi", "Kuralda olmayan belge türleri", belge_yol.name, f"{', '.join(r['taninmayan'])}. Bu belgeler hiçbir kuralla eşleşmedi; yazım farkıysa kural dosyasındaki 'Diğer Adlar' sütununa ekleyin."))
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Taburcu dosyalarındaki zorunlu belgelerin eksiksizliğini kontrol eder.")
    p.add_argument("--taburcu", type=Path, default=ORNEK / "taburcu_listesi.csv", help="Protokol No, Servis, Hekim, Taburcu Tarihi ve koşul sütunları (ör. Ameliyat, Konsültasyon: Evet / Hayır)")
    p.add_argument("--belgeler", type=Path, default=ORNEK / "belge_listesi.csv", help="Protokol No, Belge (dosyada bulunan her belge bir satır)")
    p.add_argument("--kurallar", type=Path, default=ORNEK / "belge_kurallari.csv", help="Belge, Koşul (boş / Hepsi / sütun adı), Diğer Adlar")
    p.add_argument("--sure-gun", help="Dosyanın taburcudan sonra tamamlanması gereken gün sayısı (kurum kuralınız; verilmezse kullanılmaz; örnek veride 3)")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 12.10.2026)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "taburcu_dosyasi_kontrolu.xlsx")
    a = p.parse_args(argv)
    ornek = a.taburcu == ORNEK / "taburcu_listesi.csv"
    rapor = tarih(a.tarih) if a.tarih else (date(2026, 10, 12) if ornek else date.today())
    sg = para(a.sure_gun or ("3" if ornek else None))
    if rapor is None or (a.sure_gun and (sg is None or sg < 0 or sg != int(sg))):
        print("[X] --tarih GG.AA.YYYY biçiminde; --sure-gun sıfır veya pozitif tam sayı olmalı")
        return 2
    for y in (a.taburcu, a.belgeler, a.kurallar):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.taburcu, a.belgeler, a.kurallar, a.cikti, rapor, None if sg is None else int(sg))
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    y = r["say"]
    print(f"[OK] {y['dosya']} dosya · {y['tam']} tam, {y['eksik']} eksik ({y['asildi']} süresi aşılmış) · tamlık %{r['tamlik']} · {y['eksik_belge']} eksik belge")
    for d in r["dosyalar"]:
        if d["durum"] != "Tam":
            print(f"[!] {d['protokol']} · {d['servis']}: {', '.join(d['eksik'])}" + (f" ({d['gecen']} gün)" if d["gecen"] is not None else ""))
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
