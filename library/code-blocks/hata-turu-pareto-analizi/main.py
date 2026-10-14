"""
Hata Türü Pareto Analizi — Workers / Workless kod bloğu
Kalite Kontrol › Kalite Kontrol Şefi

Ara ve final kontrol hata kayıtlarını hata türü, operasyon ve hat bazında Pareto analizine tabi tutar.
  Pay %              = kalemin hata adedi ÷ toplam hata adedi × 100
  Kümülatif %        = kalemler büyükten küçüğe sıralandığında payların birikimli toplamı
  A sınıfı           = kendisinden önceki kümülatif pay --esik (%80) değerinin altında olan kalem
  Hata oranı (DHU)   = hata adedi ÷ kontrol edilen adet × 100       (kontrol adetleri dosyası verilirse; 100 parçadaki hata sayısı)
  Uyarı              = hattın hata oranı --dhu (5) üstünde · "Diğer" grubunun payı %10'un üstünde · türü yazılmamış hata
Hata oranı hatalı parça yüzdesi değildir: bir parçada birden çok hata sayılabilir. Hata türünü kontrolör kaydı belirler.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"tarih": ("tarih", "kontrol tarihi"), "hat": ("hat", "bant", "uretim hatti", "grup"), "operasyon": ("operasyon", "islem", "operasyon adi"), "tur": ("hata turu", "hata", "hata tanimi", "hata kodu"),
         "kontrol": ("kontrol", "kontrol turu", "kontrol noktasi", "asama"), "adet": ("hata adedi", "adet", "hata sayisi")}
K_SUTUN = {"tarih": ("tarih", "kontrol tarihi"), "hat": ("hat", "bant", "uretim hatti", "grup"), "adet": ("kontrol edilen adet", "kontrol adedi", "kontrol edilen", "adet")}
YOK = "(tür yazılmamış)"
DIGER = {"diger", "diger hatalar", "muhtelif"}
SIFIR = Decimal(0)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("hat", "tur")):
        hat = metin(r["hat"])
        if not hat:
            continue
        adet = para(r["adet"]) if metin(r["adet"]) else Decimal(1)
        if adet is None or adet <= 0:
            uy.append(uyari("Yüksek", "Hata kaydı okunamadı", f"Satır {r['_satir']} · {hat}", f"Hata adedi '{metin(r['adet'])}' okunamadı. Satır alınmadı."))
            continue
        l.append({"tarih": tarih(r["tarih"]), "hat": hat, "operasyon": metin(r["operasyon"]) or "(operasyon yok)", "tur": metin(r["tur"]) or YOK, "kontrol": metin(r["kontrol"]) or "(belirtilmemiş)", "adet": adet})
    return l, uy


def oku_kontrol(yol: Path) -> tuple[dict, list[dict]]:
    uy, d = [], {}
    for r in kayitlar(yol, K_SUTUN, ("hat", "adet")):
        hat = metin(r["hat"])
        if not hat:
            continue
        adet = para(r["adet"])
        if adet is None or adet <= 0:
            uy.append(uyari("Yüksek", "Kontrol adedi okunamadı", f"Satır {r['_satir']} · {hat}", f"{yol.name}: kontrol edilen adet '{metin(r['adet'])}' okunamadı. Satır alınmadı."))
            continue
        d[katla(hat)] = d.get(katla(hat), SIFIR) + adet
    return d, uy


def _oran(pay, payda) -> Decimal | None:
    return yuvarla(Decimal(pay) * 100 / Decimal(payda)) if payda else None


def pareto(kayit: list[dict], alan: str, esik: Decimal = Decimal(80)) -> list[dict]:
    grup = {}
    for h in kayit:
        g = grup.setdefault(katla(h[alan]), {"ad": h[alan], "kayit": 0, "adet": SIFIR})
        g["kayit"], g["adet"] = g["kayit"] + 1, g["adet"] + h["adet"]
    s, toplam, kum = sorted(grup.values(), key=lambda g: (-g["adet"], katla(g["ad"]))), sum(g["adet"] for g in grup.values()), SIFIR
    for g in s:
        g["pay"], g["sinif"] = _oran(g["adet"], toplam), "A" if _oran(kum, toplam) < esik else "B"
        kum += g["adet"]
        g["kum"] = _oran(kum, toplam)
    return s


def analiz(hatalar: list[dict], kontrol: dict | None = None, esik: Decimal = Decimal(80), dhu: Decimal = Decimal(5)) -> dict:
    if not hatalar:
        raise ValueError("Analiz edilecek hata kaydı yok.")
    kontrol, uy = kontrol or {}, []
    turler, operasyonlar, hatlar = pareto(hatalar, "tur", esik), pareto(hatalar, "operasyon", esik), pareto(hatalar, "hat", esik)
    for t in turler:
        if t["ad"] == YOK:
            uy.append(uyari("Orta", "Türü yazılmamış hata", YOK, f"{t['adet']:f} hata (%{t['pay']}) türsüz kaydedilmiş."))
        elif katla(t["ad"]) in DIGER and t["pay"] > 10:
            uy.append(uyari("Orta", "'Diğer' grubu büyük", t["ad"], f"Hataların %{t['pay']} kadarı ayrıştırılmamış ({t['adet']:f} adet). Bu grubu alt türlere bölün."))
    for h in hatlar:
        en = pareto([x for x in hatalar if katla(x["hat"]) == katla(h["ad"])], "tur", esik)[0]
        h["en_tur"], h["en_pay"], h["kontrol"] = en["ad"], en["pay"], kontrol.get(katla(h["ad"]))
        h["dhu"] = _oran(h["adet"], h["kontrol"]) if h["kontrol"] else None
        if h["dhu"] is not None and h["dhu"] > dhu:
            uy.append(uyari("Orta", "Hata oranı yüksek", h["ad"], f"{h['kontrol']:f} kontrol edilen adette {h['adet']:f} hata (100 parçada {h['dhu']}; sınır {dhu}). En sık hata: {h['en_tur']} (%{h['en_pay']})."))
        elif kontrol and h["kontrol"] is None:
            uy.append(uyari("Bilgi", "Kontrol adedi yok", h["ad"], "Kontrol adetleri dosyasında bu hat yok; hata oranı hesaplanmadı."))
    toplam, top_kontrol = sum((h["adet"] for h in hatalar), SIFIR), sum((h["kontrol"] for h in hatlar if h["kontrol"]), SIFIR)
    matris = [{"hat": h["ad"], "hucre": [sum((x["adet"] for x in hatalar if katla(x["hat"]) == katla(h["ad"]) and katla(x["tur"]) == katla(t["ad"])), SIFIR) for t in turler], "toplam": h["adet"]} for h in sorted(hatlar, key=lambda h: katla(h["ad"]))]
    a = [t for t in turler if t["sinif"] == "A"]
    return {"hatalar": hatalar, "turler": turler, "operasyonlar": operasyonlar, "hatlar": hatlar, "kontroller": sorted(pareto(hatalar, "kontrol", esik), key=lambda k: katla(k["ad"])), "matris": matris,
            "ozet": {"kayit": len(hatalar), "adet": toplam, "tur": len(turler), "a_tur": len(a), "a_pay": a[-1]["kum"], "kontrol": top_kontrol if kontrol else None,
                     "dhu": _oran(sum((h["adet"] for h in hatlar if h["kontrol"]), SIFIR), top_kontrol) if top_kontrol else None}, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, esik: Decimal, dhu: Decimal) -> None:
    wb = Workbook()
    for i, (baslik, alan, ad) in enumerate((("Hata Türü Pareto", "turler", "Hata Türü"), ("Operasyon Pareto", "operasyonlar", "Operasyon"), ("Hat Pareto", "hatlar", "Hat"))):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title, hat = baslik, alan == "hatlar"
        _baslik(ws, ["Sıra", ad, "Sınıf", "Hata Adedi", "Pay %", "Kümülatif %", "Kayıt"] + (["Kontrol Edilen", "Hata Oranı (100 parçada)", "En Sık Hata", "Hattaki Payı %"] if hat else []), (6, 32, 7, 11, 9, 12, 8, 12, 12, 28, 10))
        for n, g in enumerate(r[alan], 1):
            _satir(ws, [n, g["ad"], g["sinif"], g["adet"], g["pay"], g["kum"], g["kayit"]] + ([g["kontrol"], g["dhu"], g["en_tur"], g["en_pay"]] if hat else []), para_s=(5, 6, 9, 11),
                   dolgu={3: KIRMIZI if g["sinif"] == "A" else None, 9: KIRMIZI if hat and g["dhu"] is not None and g["dhu"] > dhu else None})
        ws.append([])
        ws.append([f"A sınıfı: toplam hatanın ilk %{esik} kadarını oluşturan kalemler."])
    ws = wb.create_sheet("Hat × Hata Türü")
    _baslik(ws, ["Hat \\ Hata Türü"] + [t["ad"] for t in r["turler"]] + ["Toplam"], (18,) + (14,) * (len(r["turler"]) + 1), "B2")
    for m in r["matris"]:
        _satir(ws, [m["hat"]] + [h or None for h in m["hucre"]] + [m["toplam"]])
    ws = wb.create_sheet("Özet")
    o = r["ozet"]
    _baslik(ws, ["Gösterge", "Değer"], (44, 14))
    for etiket, deger, para_mi in (("Hata kaydı", o["kayit"], False), ("Toplam hata adedi", o["adet"], False), ("Farklı hata türü", o["tur"], False), ("A sınıfı hata türü sayısı", o["a_tur"], False), ("A sınıfı türlerin toplamdaki payı %", o["a_pay"], True),
                                   ("Kontrol edilen adet", o["kontrol"], False), ("Hata oranı (100 parçada hata)", o["dhu"], True)):
        _satir(ws, [etiket, deger], para_s=(2,) if para_mi else ())
    ws.append([])
    _baslik(ws, ["Kontrol", "Hata Adedi", "Pay %"], dondur="A2")
    for k in r["kontroller"]:
        _satir(ws, [k["ad"], k["adet"], k["pay"]], para_s=(3,))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, kontrol_yol: Path | None = None, esik: Decimal = Decimal(80), dhu: Decimal = Decimal(5)) -> dict:
    hatalar, uy = oku(yol)
    kontrol, uy2 = oku_kontrol(kontrol_yol) if kontrol_yol else ({}, [])
    r = analiz(hatalar, kontrol, esik, dhu)
    r["uyarilar"] = sirala(uy + uy2 + r["uyarilar"])
    rapor_yaz(cikti, r, esik, dhu)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Ara ve final kontrol hata kayıtlarını hata türü, operasyon ve hat bazında Pareto analizine tabi tutar.")
    p.add_argument("--hatalar", type=Path, help="Tarih, Hat, Operasyon, Hata Türü, Kontrol (Ara/Final), Hata Adedi")
    p.add_argument("--kontrol", type=Path, help="İsteğe bağlı: Tarih, Hat, Kontrol Edilen Adet; verilirse hat bazında hata oranı hesaplanır")
    p.add_argument("--esik", default="80", help="A sınıfı için kümülatif pay sınırı, yüzde (varsayılan 80)")
    p.add_argument("--dhu", default="5", help="100 parçadaki hata sayısı bu değeri aşan hat uyarılır (varsayılan 5)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "hata_turu_pareto.xlsx")
    a = p.parse_args(argv)
    ornek = a.hatalar is None
    yol, k_yol = (ORNEK / "hata_kayitlari.csv", a.kontrol or ORNEK / "kontrol_adetleri.csv") if ornek else (a.hatalar, a.kontrol)
    esik, dhu = para(a.esik), para(a.dhu)
    if esik is None or not 0 < esik <= 100 or dhu is None or dhu < 0:
        print("[X] --esik 0 ile 100 arasında; --dhu eksi olamaz")
        return 2
    for y in (yol, k_yol):
        if y is not None and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(yol, a.cikti, k_yol, esik, dhu)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {o['adet']:f} hata · {o['tur']} hata türü · A sınıfı {o['a_tur']} tür toplamın %{o['a_pay']} kadarı" + (f" · hata oranı 100 parçada {o['dhu']}" if o["dhu"] is not None else ""))
    for t in r["turler"]:
        if t["sinif"] == "A":
            print(f"     {t['ad']}: {t['adet']:f} adet · %{t['pay']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
