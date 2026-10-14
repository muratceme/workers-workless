import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def h(hat, operasyon, tur, adet, kontrol="Ara"):
    return {"tarih": date(2026, 10, 5), "hat": hat, "operasyon": operasyon, "tur": tur or main.YOK, "kontrol": kontrol, "adet": D(adet)}


def veri():
    return [h("H1", "Op1", "Dikiş atlaması", 10), h("H1", "Op2", "Leke", 5), h("h2", "OP1", "dikiş atlaması", 6), h("H2", "Op3", "Ölçü hatası", 4, "Final"), h("H2", "Op3", "", 2, "Final")]


def k(l):
    return [(g["ad"], g["sinif"], g["adet"], g["pay"], g["kum"], g["kayit"]) for g in l]


class AnalizTesti(unittest.TestCase):
    def test_pareto(self):
        r = main.analiz(veri(), {"h1": D(300), "h2": D(100)})
        self.assertEqual(k(r["turler"]), [("Dikiş atlaması", "A", D(16), D("59.26"), D("59.26"), 2), ("Leke", "A", D(5), D("18.52"), D("77.78"), 1), ("Ölçü hatası", "A", D(4), D("14.81"), D("92.59"), 1),
                                          (main.YOK, "B", D(2), D("7.41"), D("100.00"), 1)])
        self.assertEqual([(g["ad"], g["adet"]) for g in r["operasyonlar"]], [("Op1", D(16)), ("Op3", D(6)), ("Op2", D(5))])
        self.assertEqual([(g["ad"], g["adet"], g["pay"], g["kontrol"], g["dhu"], g["en_tur"], g["en_pay"]) for g in r["hatlar"]],
                         [("H1", D(15), D("55.56"), D(300), D("5.00"), "Dikiş atlaması", D("66.67")), ("h2", D(12), D("44.44"), D(100), D("12.00"), "dikiş atlaması", D("50.00"))])
        self.assertEqual([(m["hat"], m["hucre"], m["toplam"]) for m in r["matris"]], [("H1", [D(10), D(5), D(0), D(0)], D(15)), ("h2", [D(6), D(0), D(4), D(2)], D(12))])
        self.assertEqual([(g["ad"], g["adet"], g["pay"]) for g in r["kontroller"]], [("Ara", D(21), D("77.78")), ("Final", D(6), D("22.22"))])
        self.assertEqual((r["ozet"], sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"])),
                         ({"kayit": 5, "adet": D(27), "tur": 4, "a_tur": 3, "a_pay": D("92.59"), "kontrol": D(400), "dhu": D("6.75")}, [("Orta", "Hata oranı yüksek", "h2"), ("Orta", "Türü yazılmamış hata", main.YOK)]))

    def test_secenekler(self):
        r = main.analiz(veri())
        self.assertEqual((r["ozet"]["kontrol"], r["ozet"]["dhu"], [g["dhu"] for g in r["hatlar"]], [u["tur"] for u in r["uyarilar"]]), (None, None, [None, None], ["Türü yazılmamış hata"]))
        r = main.analiz(veri(), {"h1": D(300)}, esik=D("59.26"), dhu=D("4.99"))
        self.assertEqual(([g["sinif"] for g in r["turler"]], r["ozet"]["dhu"], sorted((u["tur"], u["kim"]) for u in r["uyarilar"])),
                         (["A", "B", "B", "B"], D("5.00"), [("Hata oranı yüksek", "H1"), ("Kontrol adedi yok", "h2"), ("Türü yazılmamış hata", main.YOK)]))
        r = main.analiz([h("H", "O", "DİĞER", 2), h("H", "O", "B", 8), h("H", "O", "A", 8)])
        self.assertEqual(([g["ad"] for g in r["turler"]], [u["tur"] for u in r["uyarilar"]]), (["A", "B", "DİĞER"], ["'Diğer' grubu büyük"]))
        with self.assertRaises(ValueError):
            main.analiz([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "hata_kayitlari.csv", cikti, ORNEK / "kontrol_adetleri.csv")
            self.assertEqual(k(r["turler"]), [("Dikiş atlaması", "A", D(34), D("29.31"), D("29.31"), 3), ("Kol potluğu", "A", D(24), D("20.69"), D("50.00"), 2), ("İplik temizliği", "A", D(18), D("15.52"), D("65.52"), 2),
                                              ("Yaka potluğu", "A", D(13), D("11.21"), D("76.72"), 2), ("Leke", "A", D(12), D("10.34"), D("87.07"), 2), ("Dikiş kayması", "B", D(10), D("8.62"), D("95.69"), 2),
                                              ("Ölçü hatası", "B", D(3), D("2.59"), D("98.28"), 1), ("Diğer", "B", D(2), D("1.72"), D("100.00"), 1)])
            self.assertEqual([(g["ad"], g["adet"]) for g in r["operasyonlar"]], [("Yan dikiş", D(40)), ("Etek reçme", D(24)), ("Kol takma", D(24)), ("Ütü paket", D(15)), ("Yaka takma", D(13))])
            self.assertEqual([(g["ad"], g["adet"], g["kontrol"], g["dhu"], g["en_tur"]) for g in r["hatlar"]],
                             [("Hat 1", D(46), D(820), D("5.61"), "Dikiş atlaması"), ("Hat 2", D(43), D(770), D("5.58"), "Kol potluğu"), ("Hat 3", D(27), D(610), D("4.43"), "İplik temizliği")])
            self.assertEqual((r["ozet"], [(g["ad"], g["adet"]) for g in r["kontroller"]]), ({"kayit": 15, "adet": D(116), "tur": 8, "a_tur": 5, "a_pay": D("87.07"), "kontrol": D(2200), "dhu": D("5.27")}, [("Ara", D(77)), ("Final", D(39))]))
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]], [("Orta", "Hata oranı yüksek", "Hat 1"), ("Orta", "Hata oranı yüksek", "Hat 2")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Hata Türü Pareto", "Operasyon Pareto", "Hat Pareto", "Hat × Hata Türü", "Özet", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, ky = Path(tmp) / "h.csv", Path(tmp) / "k.csv"
            y.write_text("Kontrol Tarihi;Bant;İşlem;Hata;Aşama;Adet\n05.10.2026;B1;Yaka;Potluk;Final;1.200\n;B1;;;;\n;B2;X;Y;;0\n;B2;X;Y;;z\n;;X;Y;;1\n", encoding="utf-8")
            ky.write_text("Bant;Kontrol Adedi\nB1;1.000\nb1;500\nB2;0\n;5\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["tarih"], x["hat"], x["operasyon"], x["tur"], x["kontrol"], x["adet"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([(date(2026, 10, 5), "B1", "Yaka", "Potluk", "Final", D(1200)), (None, "B1", "(operasyon yok)", main.YOK, "(belirtilmemiş)", D(1))],
                              [("Hata kaydı okunamadı", "Satır 4 · B2"), ("Hata kaydı okunamadı", "Satır 5 · B2")]))
            self.assertEqual((main.oku_kontrol(ky)[0], [u["tur"] for u in main.oku_kontrol(ky)[1]]), ({"b1": D(1500)}, ["Kontrol adedi okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--hatalar", str(y), "--cikti", c]
            self.assertEqual((main.main(arg), main.main(arg + ["--kontrol", str(ky), "--esik", "70", "--dhu", "3"])), (0, 0))
            self.assertEqual([main.main(arg + v) for v in (["--esik", "0"], ["--dhu", "-1"], ["--dhu", "x"])], [2] * 3)
            self.assertEqual((main.main(["--hatalar", str(Path(tmp) / "yok.csv")]), main.main(arg + ["--kontrol", str(Path(tmp) / "yok.csv")])), (1, 1))
            for icerik in ("Hat;Adet\nA;5\n", "Hat;Hata Türü\n;B\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
