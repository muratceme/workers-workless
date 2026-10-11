"""
Suistimal Risk Göstergeleri — Workers / Workless kod bloğu
Hasar Yönetimi › Hasar Uzmanı

Hasar dosyalarını dokuz göstergeyle tarar; tetiklenen göstergelerin puanlarını toplayıp inceleme önceliği çıkarır.
  YAKIN_BASLANGIC  hasar, poliçe başlangıcından sonraki ilk N gün içinde          YAKIN_BITIS  bitişten önceki son N gün içinde
  GEC_IHBAR        ihbar, hasardan N günden fazla sonra                           SIK_HASAR    aynı sigortalının son 365 günde N ve üzeri hasarı
  ZEYIL_SONRASI    teminat artırım zeylinden sonraki ilk N gün içinde hasar       BEDELE_YAKIN talep ≥ sigorta bedelinin %N'i
  EKSPER_FARKI     talep, eksper tespitinin %N'inden fazla üzerinde               TUTANAK_YOK  resmi tutanak yok
  BEYAN_TUTARSIZ   beyan edilen önceki hasar sayısı, dosyadaki kayıtlı önceki hasar sayısından az
  Puan = tetiklenen göstergelerin puanları toplamı (en çok 100);  düzey = --yuksek (50) ve --orta (25) eşiklerine göre
Puanlar ve eşikler örnek varsayılanlardır; göstergeler dosyasıyla değiştirilir. Gösterge kanıt değildir: yüksek puan suistimal olduğunu
göstermez, yalnız dosyanın ayrıntılı incelenmesini önerir. Ret veya ihbar kararı yalnız bu çıktıya dayandırılamaz.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"no": ("hasar no", "dosya no", "hasar dosya no"), "police": ("police no",), "sig_no": ("sigortali no", "musteri no", "tckn vkn", "kimlik no"),
         "sigortali": ("sigortali", "sigortali adi", "musteri"), "tur": ("hasar turu", "tur", "brans"), "bas": ("police baslangic", "baslangic", "baslangic tarihi"),
         "bit": ("police bitis", "bitis", "bitis tarihi"), "zeyil": ("zeyil tarihi", "teminat artirim tarihi"), "hasar": ("hasar tarihi",), "ihbar": ("ihbar tarihi", "bildirim tarihi"),
         "bedel": ("sigorta bedeli", "bedel", "teminat limiti"), "talep": ("talep tutari", "talep edilen tutar", "talep"), "eksper": ("eksper tutari", "eksper tespiti", "ekspertiz tutari"),
         "beyan": ("beyan edilen onceki hasar", "onceki hasar beyani", "beyan edilen hasar sayisi"), "tutanak": ("resmi tutanak", "tutanak", "zabit")}
# kod: (ad, varsayılan puan, varsayılan eşik, eşik birimi)
GOSTERGE = {"YAKIN_BASLANGIC": ("Poliçe başlangıcına yakın hasar", 20, 30, "gün"), "YAKIN_BITIS": ("Poliçe bitişine yakın hasar", 10, 15, "gün"),
            "GEC_IHBAR": ("Geç ihbar", 10, 10, "gün"), "SIK_HASAR": ("Son 365 günde sık hasar", 20, 3, "adet"),
            "ZEYIL_SONRASI": ("Teminat artırımından kısa süre sonra hasar", 20, 30, "gün"), "BEDELE_YAKIN": ("Talep sigorta bedeline çok yakın", 10, 90, "%"),
            "EKSPER_FARKI": ("Talep eksper tespitinin çok üzerinde", 15, 30, "%"), "BEYAN_TUTARSIZ": ("Önceki hasar beyanı kayıtlarla tutarsız", 15, None, ""),
            "TUTANAK_YOK": ("Resmi tutanak yok", 10, None, "")}
YOK = ("yok", "hayir", "h", "0")


def varsayilan() -> dict:
    return {k: {"ad": v[0], "puan": Decimal(v[1]), "esik": Decimal(v[2]) if v[2] is not None else None, "birim": v[3]} for k, v in GOSTERGE.items()}


def ayar_oku(yol: Path) -> tuple[dict, list[dict]]:
    ayar, uy = varsayilan(), []
    for r in kayitlar(yol, {"kod": ("kod", "gosterge"), "puan": ("puan", "agirlik"), "esik": ("esik", "sinir")}, ("kod",)):
        kod = metin(r["kod"]).upper()
        if not kod:
            continue
        if kod not in ayar:
            uy.append(uyari("Orta", "Gösterge kodu tanınmadı", kod, f"{yol.name} satır {r['_satir']}: geçerli kodlar {', '.join(GOSTERGE)}."))
            continue
        for alan in ("puan", "esik"):
            if metin(r[alan]):
                v = para(r[alan])
                if v is None or v < 0 or (alan == "esik" and ayar[kod]["esik"] is None):
                    uy.append(uyari("Orta", "Gösterge ayarı okunamadı", kod, f"{yol.name} satır {r['_satir']}: {alan} '{metin(r[alan])}'. Varsayılan kullanıldı."))
                else:
                    ayar[kod][alan] = v
    return ayar, uy


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("no", "hasar")):
        no = metin(r["no"])
        if not no:
            continue
        kim = f"{no} (satır {r['_satir']})"
        if no in gorulen:
            uy.append(uyari("Yüksek", "Mükerrer hasar dosyası", kim, "Hasar numarası tekrar ediyor; ilk satır kullanıldı."))
            continue
        h = {"no": no, "police": metin(r["police"]), "sigortali": metin(r["sigortali"]), "tur": metin(r["tur"]), "anahtar": katla(r["sig_no"]) or katla(r["sigortali"])}
        for a, ad in (("hasar", "Hasar tarihi"), ("ihbar", "İhbar tarihi"), ("bas", "Poliçe başlangıç"), ("bit", "Poliçe bitiş"), ("zeyil", "Zeyil tarihi")):
            h[a] = tarih(r[a])
            if metin(r[a]) and not h[a]:
                uy.append(uyari("Orta", "Tarih okunamadı", kim, f"{ad} '{metin(r[a])}'. İlgili gösterge değerlendirilemedi."))
        if not h["hasar"]:
            uy.append(uyari("Yüksek", "Hasar tarihi yok", kim, "Hasar tarihi boş veya okunamadı; dosya taranmadı."))
            continue
        for a, ad in (("bedel", "Sigorta bedeli"), ("talep", "Talep tutarı"), ("eksper", "Eksper tutarı"), ("beyan", "Beyan edilen önceki hasar")):
            h[a] = para(r[a])
            if metin(r[a]) and h[a] is None:
                uy.append(uyari("Orta", "Tutar okunamadı", kim, f"{ad} '{metin(r[a])}'. İlgili gösterge değerlendirilemedi."))
        h["tutanak"] = None if not metin(r["tutanak"]) else katla(r["tutanak"]) not in YOK
        if h["bas"] and h["bit"] and not h["bas"] <= h["hasar"] <= h["bit"]:
            uy.append(uyari("Yüksek", "Hasar tarihi poliçe süresi dışında", kim, f"Hasar {h['hasar']:%d.%m.%Y}, poliçe {h['bas']:%d.%m.%Y} – {h['bit']:%d.%m.%Y}. Teminat kontrolü gerekir."))
        gorulen.add(no)
        l.append(h)
    return l, uy


def tara(hasarlar: list[dict], ayar: dict | None = None, yuksek: Decimal = Decimal(50), orta: Decimal = Decimal(25)) -> dict:
    ayar = ayar or varsayilan()
    gecmis = defaultdict(list)
    for h in hasarlar:
        if h["anahtar"]:
            gecmis[h["anahtar"]].append(h["hasar"])
    sayac = {k: {"tetik": 0, "yok": 0} for k in ayar}
    for h in hasarlar:
        e = {k: v["esik"] for k, v in ayar.items()}
        t = h["hasar"]
        g = gecmis.get(h["anahtar"], [])
        son365, onceki = sum(0 <= (t - x).days <= 365 for x in g), sum(x < t for x in g)
        # her gösterge: None = veri yok (değerlendirilemedi), yoksa (tetiklendi mi, açıklama)
        s = {
            "YAKIN_BASLANGIC": None if not h["bas"] else (0 <= (t - h["bas"]).days <= e["YAKIN_BASLANGIC"], f"Poliçe başlangıcından {(t - h['bas']).days} gün sonra."),
            "YAKIN_BITIS": None if not h["bit"] else (0 <= (h["bit"] - t).days <= e["YAKIN_BITIS"], f"Poliçe bitişine {(h['bit'] - t).days} gün kala."),
            "GEC_IHBAR": None if not h["ihbar"] else ((h["ihbar"] - t).days > e["GEC_IHBAR"], f"İhbar hasardan {(h['ihbar'] - t).days} gün sonra."),
            "SIK_HASAR": None if not h["anahtar"] else (son365 >= e["SIK_HASAR"], f"Sigortalının bu hasar dahil son 365 günde {son365} hasarı var."),
            "ZEYIL_SONRASI": None if not h["zeyil"] else (0 <= (t - h["zeyil"]).days <= e["ZEYIL_SONRASI"], f"Teminat artırım zeylinden {(t - h['zeyil']).days} gün sonra."),
            "BEDELE_YAKIN": None if not h["bedel"] or h["talep"] is None else
            (h["talep"] / h["bedel"] * 100 >= e["BEDELE_YAKIN"], f"Talep, sigorta bedelinin %{yuvarla(h['talep'] / h['bedel'] * 100)} kadarı."),
            "EKSPER_FARKI": None if not h["eksper"] or h["talep"] is None else
            ((h["talep"] - h["eksper"]) / h["eksper"] * 100 > e["EKSPER_FARKI"], f"Talep, eksper tespitinin %{yuvarla((h['talep'] - h['eksper']) / h['eksper'] * 100)} üzerinde."),
            "BEYAN_TUTARSIZ": None if h["beyan"] is None or not h["anahtar"] else (h["beyan"] < onceki, f"Beyan edilen önceki hasar {int(h['beyan'])}, dosyadaki kayıtlı önceki hasar {onceki}."),
            "TUTANAK_YOK": None if h["tutanak"] is None else (not h["tutanak"], "Resmi tutanak yok olarak işaretlenmiş."),
        }
        h["tetiklenen"], h["eksik"], puan = [], [], Decimal(0)
        for kod, v in s.items():
            if ayar[kod]["puan"] == 0:
                continue
            if v is None:
                h["eksik"].append(kod)
                sayac[kod]["yok"] += 1
            elif v[0]:
                h["tetiklenen"].append((kod, ayar[kod]["ad"], ayar[kod]["puan"], v[1]))
                puan += ayar[kod]["puan"]
                sayac[kod]["tetik"] += 1
        h["puan"] = min(puan, Decimal(100))
        h["duzey"] = "Yüksek" if h["puan"] >= yuksek else "Orta" if h["puan"] >= orta else "Düşük"
    hasarlar.sort(key=lambda h: (-h["puan"], h["no"]))
    return {"hasarlar": hasarlar, "ayar": ayar, "sayac": sayac, "esik": (yuksek, orta)}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "İnceleme Önceliği"
    _baslik(ws, ["Hasar No", "Poliçe No", "Sigortalı", "Hasar Türü", "Hasar Tarihi", "Talep Tutarı", "Puan", "Düzey", "Gösterge Sayısı", "Tetiklenen Göstergeler", "Değerlendirilemeyen",
                 "İnceleme Sonucu (siz doldurun)"], (14, 14, 28, 18, 12, 16, 8, 10, 9, 60, 30, 30))
    renk = {"Yüksek": KIRMIZI, "Orta": SARI, "Düşük": None}
    for h in r["hasarlar"]:
        _satir(ws, [h["no"], h["police"], h["sigortali"], h["tur"], h["hasar"], h["talep"], h["puan"], h["duzey"], len(h["tetiklenen"]), "; ".join(t[1] for t in h["tetiklenen"]),
                    ", ".join(h["eksik"]), ""], para_s=(6,), tarih_s=(5,), dolgu={8: renk[h["duzey"]]}, sar=(10, 11))
    ws.append([])
    ws.append(["Gösterge kanıt değildir. Yüksek puan suistimal olduğunu göstermez; yalnız dosyanın ayrıntılı incelenmesini önerir."])
    ws = wb.create_sheet("Gösterge Ayrıntısı")
    _baslik(ws, ["Hasar No", "Sigortalı", "Düzey", "Gösterge", "Puan", "Açıklama"], (14, 28, 10, 44, 8, 80))
    for h in r["hasarlar"]:
        for kod, ad, puan, ac in h["tetiklenen"]:
            _satir(ws, [h["no"], h["sigortali"], h["duzey"], ad, puan, ac], dolgu={3: renk[h["duzey"]]}, sar=(6,))
    ws = wb.create_sheet("Gösterge Özeti")
    _baslik(ws, ["Kod", "Gösterge", "Puan", "Eşik", "Birim", "Tetiklenen Dosya", "Değerlendirilemeyen Dosya"], (20, 46, 8, 8, 8, 12, 14))
    for kod, a in r["ayar"].items():
        _satir(ws, [kod, a["ad"], a["puan"], a["esik"], a["birim"], r["sayac"][kod]["tetik"], r["sayac"][kod]["yok"]])
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, ayar_yol: Path | None = None, **esik) -> dict:
    ayar, uy1 = ayar_oku(ayar_yol) if ayar_yol else (varsayilan(), [])
    hasarlar, uy2 = oku(yol)
    if not hasarlar:
        raise ValueError(f"{yol.name}: taranacak hasar dosyası yok.")
    r = tara(hasarlar, ayar, **esik)
    r["uyarilar"] = sirala(uy1 + uy2)
    rapor_yaz(cikti, r, uy1 + uy2)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Hasar dosyalarını suistimal risk göstergeleriyle tarar ve inceleme önceliği çıkarır.")
    p.add_argument("--hasarlar", type=Path, default=ORNEK / "hasarlar.csv",
                   help="Hasar No, Poliçe No, Sigortalı No, Sigortalı, Hasar Türü, Poliçe Başlangıç, Poliçe Bitiş, Zeyil Tarihi, Hasar Tarihi, İhbar Tarihi, Sigorta Bedeli, Talep Tutarı, ...")
    p.add_argument("--gostergeler", type=Path, help="İsteğe bağlı ayar dosyası: Kod, Puan, Eşik (puan 0 = gösterge kapalı)")
    p.add_argument("--yuksek", default="50", help="Yüksek düzey için asgari puan (varsayılan 50)")
    p.add_argument("--orta", default="25", help="Orta düzey için asgari puan (varsayılan 25)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "suistimal_gostergeleri.xlsx")
    a = p.parse_args(argv)
    for y in (a.hasarlar, a.gostergeler):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    yuksek, orta = para(a.yuksek), para(a.orta)
    if yuksek is None or orta is None or not 0 < orta < yuksek <= 100:
        print("[X] 0 < --orta < --yuksek ≤ 100 olmalı")
        return 2
    try:
        r = calistir(a.hasarlar, a.cikti, a.gostergeler, yuksek=yuksek, orta=orta)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    say = lambda d: sum(h["duzey"] == d for h in r["hasarlar"])   # noqa: E731
    print(f"[OK] {len(r['hasarlar'])} hasar dosyası · {say('Yüksek')} yüksek · {say('Orta')} orta · {say('Düşük')} düşük öncelik")
    for h in r["hasarlar"]:
        if h["duzey"] == "Yüksek":
            print(f"[!] {h['no']} · puan {h['puan']}: " + "; ".join(t[1] for t in h["tetiklenen"]))
    print("[i] Gösterge kanıt değildir; yüksek puan yalnız ayrıntılı inceleme önerir.")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
