import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def hasar(no="H1", t=date(2026, 6, 15), anahtar="s1", **kw):
    h = {"no": no, "police": "P", "sigortali": "X", "tur": "", "anahtar": anahtar, "hasar": t, "ihbar": None, "bas": None, "bit": None, "zeyil": None, "bedel": None, "talep": None,
         "eksper": None, "beyan": None, "tutanak": None}
    return h | {k: (D(v) if isinstance(v, int) and not isinstance(v, bool) else v) for k, v in kw.items()}


def kodlar(h):
    return [t[0] for t in h["tetiklenen"]]


def tek(**kw):
    return main.tara([hasar(**kw)])["hasarlar"][0]


class GostergeTesti(unittest.TestCase):
    def test_tarih_gostergeleri(self):
        t = date(2026, 6, 15)
        self.assertEqual(kodlar(tek(bas=date(2026, 5, 16))), ["YAKIN_BASLANGIC"])              # 30. gün dahil
        self.assertEqual(kodlar(tek(bas=date(2026, 5, 15))), [])
        self.assertEqual(kodlar(tek(bas=date(2026, 6, 20))), [])                               # poliçe başlamadan önceki hasar gösterge değil, uyarıdır
        self.assertEqual(kodlar(tek(bit=date(2026, 6, 30))), ["YAKIN_BITIS"])
        self.assertEqual(kodlar(tek(bit=date(2026, 7, 1))), [])
        self.assertEqual(kodlar(tek(ihbar=date(2026, 6, 26))), ["GEC_IHBAR"])                  # 11 gün
        self.assertEqual(kodlar(tek(ihbar=date(2026, 6, 25))), [])
        self.assertEqual(kodlar(tek(zeyil=date(2026, 5, 20))), ["ZEYIL_SONRASI"])
        self.assertEqual(kodlar(tek(zeyil=date(2026, 6, 16))), [])
        self.assertEqual(t, tek()["hasar"])

    def test_tutar_gostergeleri(self):
        self.assertEqual(kodlar(tek(bedel=1000, talep=900)), ["BEDELE_YAKIN"])
        self.assertEqual(kodlar(tek(bedel=1000, talep=899)), [])
        self.assertEqual(kodlar(tek(eksper=1000, talep=1301)), ["EKSPER_FARKI"])
        self.assertEqual(kodlar(tek(eksper=1000, talep=1300)), [])                             # tam %30 tetiklemez
        self.assertEqual(kodlar(tek(tutanak=False)), ["TUTANAK_YOK"])
        self.assertEqual(kodlar(tek(tutanak=True)), [])
        h = tek(bedel=1000, talep=960, eksper=600, tutanak=False)
        self.assertEqual((h["puan"], h["duzey"], sorted(h["eksik"])), (D(35), "Orta", ["BEYAN_TUTARSIZ", "GEC_IHBAR", "YAKIN_BASLANGIC", "YAKIN_BITIS", "ZEYIL_SONRASI"]))
        self.assertIn("%60.00 üzerinde", h["tetiklenen"][1][3])

    def test_sik_hasar_ve_beyan(self):
        l = [hasar("H1", date(2025, 6, 1)), hasar("H2", date(2026, 1, 10), beyan=0), hasar("H3", date(2026, 5, 1), beyan=2), hasar("H4", date(2026, 6, 15), beyan=1),
             hasar("H5", date(2026, 6, 15), anahtar="s2", beyan=0), hasar("H6", date(2026, 6, 15), anahtar="")]
        d = {h["no"]: kodlar(h) for h in main.tara(l)["hasarlar"]}
        self.assertEqual(d["H2"], ["BEYAN_TUTARSIZ"])                                           # 1 kayıtlı önceki hasar, beyan 0
        self.assertEqual(d["H3"], ["SIK_HASAR"])                                                # son 365 günde H1, H2, H3
        self.assertEqual(d["H4"], ["SIK_HASAR", "BEYAN_TUTARSIZ"])                              # H1 365 günden eski: H2, H3, H4; beyan 1 < 3
        self.assertEqual((d["H1"], d["H5"], d["H6"]), ([], [], []))

    def test_puan_duzey_ve_ayar(self):
        tam = dict(bas=date(2026, 6, 1), bit=date(2026, 6, 20), ihbar=date(2026, 7, 15), zeyil=date(2026, 6, 10), bedel=100, talep=100, eksper=50, beyan=0, tutanak=False)
        l = [hasar("H0", date(2026, 1, 1)), hasar("H1", date(2026, 3, 1)), hasar("H2", **tam)]
        h = next(x for x in main.tara(l)["hasarlar"] if x["no"] == "H2")
        self.assertEqual((len(h["tetiklenen"]), h["puan"], h["duzey"]), (9, D(100), "Yüksek"))   # 130 puan 100 ile sınırlanır
        ayar = main.varsayilan()
        ayar["TUTANAK_YOK"]["puan"], ayar["BEDELE_YAKIN"]["esik"] = D(0), D(99)
        r = main.tara([hasar(bedel=1000, talep=980, tutanak=False)], ayar)
        self.assertEqual((r["hasarlar"][0]["puan"], r["sayac"]["TUTANAK_YOK"]), (D(0), {"tetik": 0, "yok": 0}))
        self.assertEqual([main.tara([hasar(bedel=10, talep=10, tutanak=False, bas=date(2026, 6, 1))], yuksek=D(40), orta=o)["hasarlar"][0]["duzey"] for o in (D(25), D(39))], ["Yüksek"] * 2)
        self.assertEqual(main.tara([hasar(bedel=10, talep=10, tutanak=False)])["hasarlar"][0]["duzey"], "Düşük")       # 20 puan
        self.assertEqual(main.tara([hasar(bedel=10, talep=10, eksper=5)])["hasarlar"][0]["duzey"], "Orta")              # tam 25 puan


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "s.xlsx"
            r = main.calistir(ORNEK / "hasarlar.csv", cikti, ORNEK / "gostergeler.csv")
            d = {h["no"]: (h["puan"], h["duzey"]) for h in r["hasarlar"]}
            self.assertEqual(d, {"H-209": (D(65), "Yüksek"), "H-205": (D(60), "Yüksek"), "H-201": (D(55), "Yüksek"), "H-206": (D(30), "Orta"), "H-203": (D(20), "Düşük"),
                                 "H-207": (D(20), "Düşük"), "H-204": (D(15), "Düşük"), "H-208": (D(15), "Düşük"), "H-202": (D(0), "Düşük"), "H-210": (D(0), "Düşük")})
            self.assertEqual([h["no"] for h in r["hasarlar"]][:4], ["H-209", "H-205", "H-201", "H-206"])
            k = {h["no"]: sorted(kodlar(h)) for h in r["hasarlar"]}
            self.assertEqual(k["H-205"], ["BEYAN_TUTARSIZ", "EKSPER_FARKI", "GEC_IHBAR", "SIK_HASAR"])
            self.assertEqual((k["H-206"], k["H-208"]), (["BEDELE_YAKIN", "ZEYIL_SONRASI"], ["EKSPER_FARKI"]))
            self.assertEqual((r["sayac"]["ZEYIL_SONRASI"], r["sayac"]["EKSPER_FARKI"], r["uyarilar"]), ({"tetik": 1, "yok": 9}, {"tetik": 4, "yok": 1}, []))
            self.assertEqual(load_workbook(cikti).sheetnames, ["İnceleme Önceliği", "Gösterge Ayrıntısı", "Gösterge Özeti", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            h, g = Path(tmp) / "h.csv", Path(tmp) / "g.csv"
            h.write_text("Hasar No;Sigortalı;Poliçe Başlangıç;Poliçe Bitiş;Hasar Tarihi;İhbar Tarihi;Talep Tutarı;Resmi Tutanak\nH1;A;01.01.2026;01.01.2027;05.02.2027;yarın;çok;HAYIR\n"
                         "H1;A;;;01.02.2026;;;\nH2;A;;;;;;\nH3;B;;;01.03.2026;;1.000,00;evet\n", encoding="utf-8")
            l, uy = main.oku(h)
            self.assertEqual([(x["no"], x["tutanak"], x["talep"], x["anahtar"]) for x in l], [("H1", False, None, "a"), ("H3", True, D("1000.00"), "b")])
            self.assertEqual(sorted(u["tur"] for u in uy), ["Hasar tarihi poliçe süresi dışında", "Hasar tarihi yok", "Mükerrer hasar dosyası", "Tarih okunamadı", "Tutar okunamadı"])
            g.write_text("Kod;Puan;Eşik\ngec_ihbar;25;5\nTUTANAK_YOK;0;\nBILINMEYEN;5;\nSIK_HASAR;çok;\nBEYAN_TUTARSIZ;;3\n", encoding="utf-8")
            a, uy = main.ayar_oku(g)
            self.assertEqual((a["GEC_IHBAR"]["puan"], a["GEC_IHBAR"]["esik"], a["TUTANAK_YOK"]["puan"], a["SIK_HASAR"]["puan"]), (D(25), D(5), D(0), D(20)))
            self.assertEqual(sorted(u["tur"] for u in uy), ["Gösterge ayarı okunamadı", "Gösterge ayarı okunamadı", "Gösterge kodu tanınmadı"])
            c = str(Path(tmp) / "x.xlsx")
            self.assertEqual(main.main(["--hasarlar", str(h), "--gostergeler", str(g), "--cikti", c]), 0)
            self.assertEqual(main.main(["--hasarlar", str(h), "--cikti", c, "--orta", "60"]), 2)
            self.assertEqual(main.main(["--hasarlar", str(Path(tmp) / "yok.csv")]), 1)
            h.write_text("Hasar No;Hasar Tarihi\n", encoding="utf-8")
            self.assertEqual(main.main(["--hasarlar", str(h), "--cikti", c]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
