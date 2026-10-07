"""
Kıdem ve İhbar Tazminatı Hesaplama — Workers / Workless kod bloğu
İnsan Kaynakları › Bordro ve Özlük İşleri Uzmanı

İşe giriş-çıkış tarihi ve giydirilmiş brüt ücretten kıdem tazminatını (1475 s. Kanun md. 14,
tavan sınırıyla) ve ihbar tazminatını (4857 s. Kanun md. 17) hesaplar; gelir vergisi ve damga
vergisi kesintileriyle net tutarları çıkarır. Tek kişi için komut satırından, çok kişi için
listeden çalışır. İnternete bağlanmaz.

Kullanım:
    python main.py --giris 2019-03-15 --cikis 2026-09-30 --brut 60000
    python main.py --giris 2019-03-15 --cikis 2026-09-30 --brut 60000 --yan-odeme 4500 --ikramiye 60000
    python main.py --girdi cikislar.xlsx
"""
from __future__ import annotations

import argparse
import csv
import sys
from dataclasses import asdict, dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import bordro

BURASI = Path(__file__).resolve().parent
para = bordro.para

# 4857 s. Kanun md. 17: bildirim süreleri (hizmet süresine göre)
IHBAR_SURELERI = [  # (hizmet süresi bu aydan AZ ise, hafta)
    (6, 2),     # 6 aydan az: 2 hafta
    (18, 4),    # 6 ay - 1,5 yıl: 4 hafta
    (36, 6),    # 1,5 - 3 yıl: 6 hafta
    (None, 8),  # 3 yıldan fazla: 8 hafta
]


def tarih(x) -> date:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x).strip()
    for f in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    raise ValueError(f"Tarih anlaşılamadı: {x!r}")


def ay_ekle(t: date, n: int) -> date:
    """t tarihine n ay ekler; gün o ayda yoksa ayın son gününe sabitlenir (31 Ocak + 1 ay = 28/29 Şubat)."""
    y, m = divmod(t.month - 1 + n, 12)
    yil, ay = t.year + y, m + 1
    sonraki = date(yil + (ay == 12), ay % 12 + 1, 1)
    return date(yil, ay, min(t.day, (sonraki - timedelta(days=1)).day))


def hizmet_suresi(giris: date, cikis: date) -> tuple[int, int, int]:
    """Giriş ve çıkış günleri dahil hizmet süresi (yıl, ay, gün)."""
    if cikis < giris:
        raise ValueError("Çıkış tarihi girişten önce olamaz")
    son = cikis + timedelta(days=1)          # çıkış günü dahil
    aylar = (son.year - giris.year) * 12 + (son.month - giris.month)
    if ay_ekle(giris, aylar) > son:
        aylar -= 1
    gun = (son - ay_ekle(giris, aylar)).days
    return aylar // 12, aylar % 12, gun


@dataclass
class Sonuc:
    ad: str
    giris: date
    cikis: date
    hizmet: str
    giydirilmis_brut: Decimal
    kidem_tavani: Decimal
    kidem_esas_ucret: Decimal
    kidem_brut: Decimal
    kidem_damga: Decimal
    kidem_net: Decimal
    ihbar_suresi_hafta: int
    ihbar_brut: Decimal
    ihbar_gelir_vergisi: Decimal
    ihbar_damga: Decimal
    ihbar_net: Decimal
    toplam_net: Decimal
    notlar: str


def hesapla(ad: str, giris, cikis, brut, yan_odeme=0, ikramiye_yillik=0, kumulatif_matrah=0,
            kidem_hakki: bool = True, ihbar_hakki: bool = True) -> Sonuc:
    giris, cikis = tarih(giris), tarih(cikis)
    p = bordro.Parametreler(cikis.year)
    yil, ay, gun = hizmet_suresi(giris, cikis)
    toplam_ay = yil * 12 + ay
    notlar = []

    # Giydirilmiş ücret: çıplak brüt + düzenli aylık yan ödemeler (yemek, yol vb.) + yıllık ikramiye / 12
    giydirilmis = para(Decimal(str(brut)) + Decimal(str(yan_odeme)) + Decimal(str(ikramiye_yillik)) / 12)
    tavan = p.kidem_tavani(cikis)
    esas = min(giydirilmis, tavan)
    if giydirilmis > tavan:
        notlar.append(f"Giydirilmiş ücret kıdem tavanını aştı; tavan ({bordro.para(tavan)}) esas alındı")

    # Kıdem: her tam yıl için 30 günlük ücret, artan süreler oranlı (1475 md. 14); en az 1 yıl hizmet
    if kidem_hakki and yil >= 1:
        kidem = para(esas * yil + esas * ay / 12 + esas * gun / 365)
    else:
        kidem = Decimal("0.00")
        notlar.append("Kıdem tazminatı hesaplanmadı: " + ("hizmet süresi 1 yıldan az" if yil < 1 else "fesih türü kıdem hakkı doğurmuyor"))
    kidem_damga = para(kidem * p.damga)     # kıdem tazminatı GVK md. 25/7 uyarınca gelir vergisinden istisna

    # İhbar: hizmet süresine göre 2-8 hafta, giydirilmiş günlük ücret üzerinden (tavan yok)
    hafta = next(h for sinir, h in IHBAR_SURELERI if sinir is None or toplam_ay < sinir)
    if ihbar_hakki:
        ihbar = para(giydirilmis / 30 * hafta * 7)
        onceki = Decimal(str(kumulatif_matrah))
        ihbar_gv = para(p.gv_hesapla(onceki + ihbar) - p.gv_hesapla(onceki))
        ihbar_dv = para(ihbar * p.damga)
        if not kumulatif_matrah:
            notlar.append("İhbar GV, yıl içi kümülatif matrah 0 kabul edilerek hesaplandı; gerçek kümülatifi girin")
    else:
        ihbar = ihbar_gv = ihbar_dv = Decimal("0.00")
        notlar.append("İhbar tazminatı hesaplanmadı (bildirim süresine uyulmuş veya hak doğmuyor)")
    ihbar_net = ihbar - ihbar_gv - ihbar_dv
    kidem_net = kidem - kidem_damga
    return Sonuc(ad=ad, giris=giris, cikis=cikis, hizmet=f"{yil} yıl {ay} ay {gun} gün",
                 giydirilmis_brut=giydirilmis, kidem_tavani=tavan, kidem_esas_ucret=esas,
                 kidem_brut=kidem, kidem_damga=kidem_damga, kidem_net=kidem_net,
                 ihbar_suresi_hafta=hafta, ihbar_brut=ihbar, ihbar_gelir_vergisi=ihbar_gv, ihbar_damga=ihbar_dv,
                 ihbar_net=ihbar_net, toplam_net=kidem_net + ihbar_net, notlar="; ".join(notlar))


# ----------------------------------------------------------------------------
# Toplu girdi ve Excel
# ----------------------------------------------------------------------------

BASLIKLAR = {
    "ad": ("ad soyad", "ad", "personel", "sicil"),
    "giris": ("işe giriş", "ise giris", "giriş tarihi", "giris tarihi", "giriş", "giris"),
    "cikis": ("işten çıkış", "isten cikis", "çıkış tarihi", "cikis tarihi", "çıkış", "cikis"),
    "brut": ("brüt", "brut", "brüt ücret", "brut ucret", "çıplak brüt"),
    "yan_odeme": ("yan ödeme", "yan odeme", "aylık yan ödemeler", "yemek+yol"),
    "ikramiye_yillik": ("yıllık ikramiye", "yillik ikramiye", "ikramiye"),
    "kumulatif_matrah": ("kümülatif matrah", "kumulatif matrah"),
    "kidem_hakki": ("kıdem hakkı", "kidem hakki"),
    "ihbar_hakki": ("ihbar hakkı", "ihbar hakki"),
}


def kucuk(s: str) -> str:
    """Türkçe küçük harf: 'İ' → 'i', 'I' → 'ı' (Python'un lower() 'İ'yi 'i̇' yapar)."""
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def evet(x) -> bool:
    return kucuk(x if x is not None else "E") not in {"h", "hayır", "hayir", "0", "false", "yok"}


def liste_oku(yol: Path) -> list[dict]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ayirici = csv.Sniffer().sniff(metin.splitlines()[0], delimiters=",;\t").delimiter
        satirlar = list(csv.reader(metin.splitlines(), delimiter=ayirici))
    satirlar = [r for r in satirlar if any(c not in (None, "") for c in r)]
    b = [kucuk(x) for x in satirlar[0]]
    konum = {alan: next((i for i, x in enumerate(b) if x in adlar), None) for alan, adlar in BASLIKLAR.items()}
    eksik = [a for a in ("ad", "giris", "cikis", "brut") if konum[a] is None]
    if eksik:
        raise SystemExit(f"Gerekli sütunlar bulunamadı: {eksik}. Başlıklar: {satirlar[0]}")
    kayitlar = []
    for r in satirlar[1:]:
        k = {alan: r[i] for alan, i in konum.items() if i is not None}
        for sayisal in ("brut", "yan_odeme", "ikramiye_yillik", "kumulatif_matrah"):
            if sayisal in k:
                v = str(k[sayisal] or "0").replace(" ", "")
                if "," in v:
                    v = v.replace(".", "").replace(",", ".")
                k[sayisal] = Decimal(v or "0")
        k["kidem_hakki"] = evet(k.get("kidem_hakki"))
        k["ihbar_hakki"] = evet(k.get("ihbar_hakki"))
        kayitlar.append(k)
    return kayitlar


ETIKETLER = {
    "ad": "Ad Soyad", "giris": "İşe Giriş", "cikis": "İşten Çıkış", "hizmet": "Hizmet Süresi",
    "giydirilmis_brut": "Giydirilmiş Brüt", "kidem_tavani": "Kıdem Tavanı", "kidem_esas_ucret": "Kıdeme Esas Ücret",
    "kidem_brut": "Kıdem Brüt", "kidem_damga": "Kıdem Damga V.", "kidem_net": "Kıdem Net",
    "ihbar_suresi_hafta": "İhbar Süresi (Hafta)", "ihbar_brut": "İhbar Brüt", "ihbar_gelir_vergisi": "İhbar Gelir V.",
    "ihbar_damga": "İhbar Damga V.", "ihbar_net": "İhbar Net", "toplam_net": "Toplam Net", "notlar": "Notlar",
}


def rapor_yaz(sonuclar: list[Sonuc], cikti: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Tazminatlar"
    alanlar = list(ETIKETLER)
    ws.append([ETIKETLER[a] for a in alanlar])
    for h in ws[1]:
        h.fill, h.font = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
    for s in sonuclar:
        d = asdict(s)
        ws.append([float(d[a]) if isinstance(d[a], Decimal) else d[a] for a in alanlar])
    for j, a in enumerate(alanlar, 1):
        ws.column_dimensions[get_column_letter(j)].width = 60 if a == "notlar" else 16
        for h in ws[get_column_letter(j)][1:]:
            if isinstance(h.value, float):
                h.number_format = "#,##0.00"
            elif isinstance(h.value, date):
                h.number_format = "DD.MM.YYYY"
    ws.freeze_panes = "B2"

    k = wb.create_sheet("Kurallar")
    for satir in [
        ["Kıdem tazminatı", "1475 s. İş Kanunu md. 14: her tam hizmet yılı için 30 günlük giydirilmiş brüt ücret; artan ay ve günler oranlanır. En az 1 yıl hizmet gerekir."],
        ["Kıdem tavanı", "Fesih tarihinde yürürlükteki tavan (tr_parametreler.json, kaynaklı)."],
        ["Kıdem vergisi", "Gelir vergisinden istisna (GVK md. 25/7); yalnızca damga vergisi (binde 7,59)."],
        ["İhbar süresi", "4857 s. İş Kanunu md. 17: 6 aydan az 2 hafta; 6 ay-1,5 yıl 4 hafta; 1,5-3 yıl 6 hafta; 3 yıldan fazla 8 hafta."],
        ["İhbar vergisi", "Gelir vergisi (yıl içi kümülatif matraha göre) ve damga vergisi; SGK primi kesilmez."],
        ["Hak koşulları", "Kıdem: işveren feshi (haklı neden hariç), işçinin haklı feshi, askerlik, emeklilik, evlilik (kadın, 1 yıl içinde), ölüm vb. İstifa ve işverenin haklı (ahlak ve iyi niyet) feshi genellikle kıdem hakkı doğurmaz. İhbar: bildirim süresine uyulmadan yapılan fesihte karşı taraf hak kazanır."],
        ["Uyarı", "Hesaplama bilgilendirme amaçlıdır. Hak doğup doğmadığı fesih nedenine bağlıdır; uyuşmazlıkta hukuki destek alın."],
    ]:
        k.append(satir)
    k.column_dimensions["A"].width = 20
    k.column_dimensions["B"].width = 130
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description="Kıdem ve ihbar tazminatı hesabı (Türkiye).")
    ap.add_argument("--girdi", type=Path, help="Çıkış listesi (.xlsx/.csv)")
    ap.add_argument("--giris", help="İşe giriş tarihi (GG.AA.YYYY veya YYYY-AA-GG)")
    ap.add_argument("--cikis", help="İşten çıkış tarihi")
    ap.add_argument("--brut", type=Decimal, help="Son aylık çıplak brüt ücret")
    ap.add_argument("--yan-odeme", type=Decimal, default=Decimal(0), help="Düzenli aylık yan ödemeler (yemek, yol vb.) toplamı")
    ap.add_argument("--ikramiye", type=Decimal, default=Decimal(0), help="Yıllık ikramiye toplamı (12'ye bölünür)")
    ap.add_argument("--kumulatif", type=Decimal, default=Decimal(0), help="Çıkış yılındaki kümülatif GV matrahı (ihbar vergisi için)")
    ap.add_argument("--kidem-yok", action="store_true", help="Fesih türü kıdem hakkı doğurmuyor")
    ap.add_argument("--ihbar-yok", action="store_true", help="İhbar tazminatı doğmuyor")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "tazminat.xlsx")
    a = ap.parse_args(argv)

    if a.giris and a.cikis and a.brut is not None:
        sonuclar = [hesapla("Çalışan", a.giris, a.cikis, a.brut, a.yan_odeme, a.ikramiye, a.kumulatif,
                            not a.kidem_yok, not a.ihbar_yok)]
    else:
        kaynak = a.girdi or BURASI / "ornek_veri" / "cikislar.csv"
        sonuclar = [hesapla(**k) for k in liste_oku(kaynak)]
    rapor_yaz(sonuclar, a.cikti)
    for s in sonuclar:
        print(f"[OK] {s.ad}: {s.hizmet} · kıdem net {s.kidem_net:,.2f} · ihbar net {s.ihbar_net:,.2f} · toplam {s.toplam_net:,.2f}"
              .replace(",", "X").replace(".", ",").replace("X", "."))
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
