import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
KADEME = [(D(0), D(1)), (D(100000), D(2)), (D(200000), D(3))]


def a(ad, tarih, tutar):
    return {"tedarikci": ad, "tarih": tarih, "tutar": D(tutar)}


def veri():
    anl = {"t1": {"tedarikci": "T1", "yontem": "Tüm ciro", "kademe": KADEME}, "t2": {"tedarikci": "T2", "yontem": "Dilimli", "kademe": KADEME},
           "t3": {"tedarikci": "T3", "yontem": "Tüm ciro", "kademe": [(D(50000), D(2))]}}
    alim = [a("T1", date(2026, 3, 1), 150000), a("t1", date(2026, 3, 2), -10000), a("T1", date(2025, 12, 31), 5000), a("T2", date(2026, 4, 1), 250000), a("X", None, 1000)]
    return anl, alim


class HesapTesti(unittest.TestCase):
    def test_prim(self):
        r = main.hesapla(*veri(), date(2026, 1, 1), date(2026, 12, 31), date(2026, 6, 30))
        self.assertEqual([(s["tedarikci"], s["brut"], s["iade"], s["ciro"], s["alt"], s["oran"], s["prim"], s["sonraki_alt"], s["kalan"], s["ek_prim"], s["tahmin"], s["tahmin_prim"]) for s in r["tedarikciler"]],
                         [("T1", D(150000), D(10000), D(140000), D(100000), D(2), D("2800.00"), D(200000), D(60000), D("3200.00"), D("282320.44"), D("8469.61")),
                          ("T2", D(250000), D(0), D(250000), D(200000), D(3), D("4500.00"), None, None, None, D("504143.65"), D("12124.31")),
                          ("T3", D(0), D(0), D(0), None, None, D(0), D(50000), D(50000), D("1000.00"), D("0.00"), D(0))])
        self.assertEqual((r["toplam_prim"], r["aylar"], r["aylik"][("t1", "2026-03")]), (D("7300.00"), ["2026-03", "2026-04"], D(140000)))
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Alım yok", "T3"), ("Anlaşma yok", "X"), ("Dönem dışı alım", ""), ("Sonraki kademeye ulaşılabilir", "T1")])

    def test_yontem_ve_sinirlar(self):
        p = lambda c, y="Dilimli", k=KADEME: main.prim_hesapla(D(c), k, y)   # noqa: E731
        self.assertEqual([p(c) for c in (0, 99999, 100000, 150000)], [(D(0), None, None), (D("999.99"), D(0), D(1)), (D("1000.00"), D(100000), D(2)), (D("2000.00"), D(100000), D(2))])
        self.assertEqual((p(150, k=[(D(100), D(2))]), p(150, "Tüm ciro", [(D(100), D(2))]), p(-5, "Tüm ciro")), ((D("1.00"), D(100), D(2)), (D("3.00"), D(100), D(2)), (D(0), None, None)))
        anl, alim = veri()
        r = main.hesapla(anl, [a("T1", None, 185000), a("T2", None, -5)])
        self.assertEqual((sorted((u["tur"], u["kim"]) for u in r["uyarilar"] if u["kim"] in ("T1", "T2")), r["tedarikciler"][0]["tahmin"]),
                         ([("Sonraki kademeye yakın", "T1"), ("İade alımdan fazla", "T2")], None))
        with self.assertRaises(ValueError):
            main.hesapla({}, alim)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "ciro_primi_anlasmalari.csv", ORNEK / "tedarikci_alimlari.csv", cikti, date(2026, 1, 1), date(2026, 12, 31), date(2026, 9, 30))
            self.assertEqual([(s["tedarikci"], s["ciro"], s["prim"], s["tahmin_prim"]) for s in r["tedarikciler"]],
                             [("Anadolu Gıda", D("2577700.00"), D("51554.00"), D("68927.51")), ("Ege Temizlik", D("1568100.00"), D("13521.50"), D("22413.60")),
                              ("Marmara İçecek", D("1753000.00"), D("35060.00"), D("46875.09"))])
            self.assertEqual([(u["tur"], u["kim"]) for u in r["uyarilar"]], [("Anlaşma yok", "Karadeniz Kuruyemiş"), ("Dönem dışı alım", ""), ("Sonraki kademeye ulaşılabilir", "Ege Temizlik")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Prim Hesabı", "Kademeler", "Aylık Ciro", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, z = Path(tmp) / "a.csv", Path(tmp) / "b.csv"
            y.write_text("Firma;Kademe;Oran;Uygulama\nA;100.000;2;\nA;0;1;dilimli\nA;100.000;3;\nA;5;x;\nA;5;1;başka\nB;0;5;\nB;10;1;\n", encoding="utf-8")
            z.write_text("Firma;Fatura Tarihi;Matrah;İşlem Türü\nA;01.01.2026;1.000;\nA;;500;İade\nA;x;5;\nA;;y;\n", encoding="utf-8")
            t, uy = main.oku_anlasma(y)
            self.assertEqual(({k: (v["yontem"], v["kademe"]) for k, v in t.items()}, sorted(u["tur"] for u in uy)),
                             ({"a": ("Tüm ciro", [(D(0), D(1)), (D(100000), D(2))]), "b": ("Tüm ciro", [(D(0), D(5)), (D(10), D(1))])},
                              ["Kademe iki kez yazılmış", "Kademe satırı okunamadı", "Kademe satırı okunamadı", "Oran kademeyle düşüyor", "Yöntem tutarsız"]))
            l, uy = main.oku_alim(z)
            self.assertEqual(([(x["tutar"], x["tarih"]) for x in l], len(uy)), ([(D(1000), date(2026, 1, 1)), (D(-500), None)], 2))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--anlasmalar", str(y), "--alimlar", str(z), "--cikti", c]
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(arg + ["--baslangic", "01.01.2026", "--bitis", "31.12.2026", "--tarih", "01.02.2026"]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--baslangic", "02.01.2026", "--bitis", "01.01.2026"])], [2] * 2)
            self.assertEqual(main.main(arg[:2] + ["--alimlar", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Firma;Oran\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
