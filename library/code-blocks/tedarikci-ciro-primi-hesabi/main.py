"""
Tedarikçi Ciro Primi Hesabı — Workers / Workless kod bloğu
Satın Alma ve Kategori Yönetimi › Kategori Yöneticisi
Tedarikçi anlaşmalarındaki ciro primi kademelerine göre dönemde hak edilen primleri hesaplar.
  Net ciro            = dönemdeki alımlar − iadeler (KDV hariç; anlaşmadaki ciro tanımına göre siz hazırlarsınız)
  Tüm ciro yöntemi    = ulaşılan en yüksek kademenin oranı tüm net ciroya uygulanır
  Dilimli yöntem      = her dilimdeki ciro kendi kademesinin oranıyla primlendirilir
  Sonraki kademe      = kalan ciro ve o kademeye ulaşılırsa eklenecek prim
  Dönem sonu tahmini  = net ciro ÷ geçen gün × dönemin gün sayısı (--baslangic, --bitis, --tarih verilirse)
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

ANLASMA = {"tedarikci": ("tedarikci", "tedarikci adi", "firma", "cari"), "alt": ("alt sinir", "kademe alt siniri", "ciro alt siniri", "kademe", "ciro"),
           "oran": ("prim", "prim orani", "prim yuzdesi", "oran"), "yontem": ("yontem", "hesap yontemi", "uygulama")}
ALIM = {"tedarikci": ("tedarikci", "tedarikci adi", "firma", "cari"), "tarih": ("tarih", "fatura tarihi", "belge tarihi"), "tutar": ("tutar", "net tutar", "kdv haric tutar", "matrah", "alis tutari"),
        "tur": ("tur", "islem turu", "belge turu")}
SIFIR = Decimal(0)


def _yontem(x) -> str | None:
    k = katla(x)
    if not k or k in ("tum ciro", "tumu", "toplam", "tum", "geriye donuk"):
        return "Tüm ciro"
    if k in ("dilimli", "dilim", "kademeli dilim", "artan oranli"):
        return "Dilimli"
    return None


def oku_anlasma(yol: Path) -> tuple[dict, list[dict]]:
    uy, t = [], {}
    for r in kayitlar(yol, ANLASMA, ("tedarikci", "alt", "oran")):
        ad = metin(r["tedarikci"])
        if not ad:
            continue
        alt, oran, yontem = para(r["alt"]), para(r["oran"]), _yontem(r["yontem"])
        if alt is None or oran is None or alt < 0 or not 0 <= oran <= 100 or yontem is None:
            uy.append(uyari("Yüksek", "Kademe satırı okunamadı", ad, f"Satır {r['_satir']}: alt sınır, prim oranı (0-100) ya da yöntem (Tüm ciro / Dilimli) okunamadı. Kademe alınmadı."))
            continue
        a = t.setdefault(katla(ad), {"tedarikci": ad, "yontem": yontem, "kademe": []})
        if yontem != a["yontem"] and metin(r["yontem"]):
            uy.append(uyari("Orta", "Yöntem tutarsız", ad, f"Satır {r['_satir']}: kademelerde farklı yöntem yazılmış; ilk satırdaki '{a['yontem']}' kullanıldı."))
        if any(k[0] == alt for k in a["kademe"]):
            uy.append(uyari("Yüksek", "Kademe iki kez yazılmış", ad, f"Satır {r['_satir']}: {alt:,.2f} alt sınırı ikinci kez geçiyor. İkinci satır alınmadı."))
            continue
        a["kademe"].append((alt, oran))
    for a in t.values():
        a["kademe"].sort()
        if any(b[1] < o[1] for o, b in zip(a["kademe"], a["kademe"][1:])):
            uy.append(uyari("Orta", "Oran kademeyle düşüyor", a["tedarikci"], "Üst kademenin oranı alt kademeden düşük; anlaşmayı kontrol edin."))
    return t, uy


def oku_alim(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, ALIM, ("tedarikci", "tutar")):
        ad = metin(r["tedarikci"])
        if not ad:
            continue
        tutar, t = para(r["tutar"]), tarih(r["tarih"]) if metin(r["tarih"]) else None
        if tutar is None or (metin(r["tarih"]) and t is None):
            uy.append(uyari("Yüksek", "Alım satırı okunamadı", ad, f"Satır {r['_satir']}: tutar ya da tarih okunamadı. Satır alınmadı."))
            continue
        if "iade" in katla(r["tur"]) and tutar > 0:
            tutar = -tutar
        l.append({"tedarikci": ad, "tarih": t, "tutar": tutar})
    return l, uy


def prim_hesapla(ciro: Decimal, kademe: list[tuple[Decimal, Decimal]], yontem: str) -> tuple[Decimal, Decimal | None, Decimal | None]:
    """(prim, ulaşılan kademe alt sınırı, oranı)."""
    ulasilan = [k for k in kademe if k[0] <= ciro and ciro > 0]
    if not ulasilan:
        return SIFIR, None, None
    alt, oran = ulasilan[-1]
    if yontem == "Tüm ciro":
        return yuvarla(ciro * oran / 100), alt, oran
    prim = SIFIR
    for i, (a, o) in enumerate(ulasilan):
        ust = kademe[i + 1][0] if i + 1 < len(kademe) else None
        prim += (min(ciro, ust) if ust is not None else ciro) * o / 100 - a * o / 100
    return yuvarla(prim), alt, oran


def hesapla(anlasmalar: dict, alimlar: list[dict], baslangic: date | None = None, bitis: date | None = None, gun: date | None = None) -> dict:
    if not anlasmalar:
        raise ValueError("Ciro primi anlaşması yok.")
    uy, ciro, iade, aylik, adlar, disarida = [], defaultdict(lambda: SIFIR), defaultdict(lambda: SIFIR), defaultdict(lambda: SIFIR), {}, 0
    for a in alimlar:
        if a["tarih"] and ((baslangic and a["tarih"] < baslangic) or (bitis and a["tarih"] > bitis)):
            disarida += 1
            continue
        k = katla(a["tedarikci"])
        adlar.setdefault(k, a["tedarikci"])
        ciro[k] += a["tutar"]
        if a["tutar"] < 0:
            iade[k] -= a["tutar"]
        if a["tarih"]:
            aylik[(k, a["tarih"].strftime("%Y-%m"))] += a["tutar"]
    if disarida:
        uy.append(uyari("Bilgi", "Dönem dışı alım", "", f"{disarida} alım satırı anlaşma dönemi dışında kaldığı için hesaba girmedi."))
    oran_gun = None
    if baslangic and bitis and gun and baslangic <= gun < bitis:
        oran_gun = (Decimal((bitis - baslangic).days + 1), Decimal((gun - baslangic).days + 1))
    sonuc = []
    for k, a in sorted(anlasmalar.items()):
        net = ciro.get(k, SIFIR)
        prim, alt, oran = prim_hesapla(net, a["kademe"], a["yontem"])
        sonraki = next((x for x in a["kademe"] if x[0] > net), None)
        s = {"tedarikci": a["tedarikci"], "yontem": a["yontem"], "brut": net + iade.get(k, SIFIR), "iade": iade.get(k, SIFIR), "ciro": net, "alt": alt, "oran": oran, "prim": prim,
             "sonraki_alt": sonraki[0] if sonraki else None, "sonraki_oran": sonraki[1] if sonraki else None, "kalan": sonraki[0] - net if sonraki else None,
             "ek_prim": prim_hesapla(sonraki[0], a["kademe"], a["yontem"])[0] - prim if sonraki else None, "tahmin": None, "tahmin_prim": None}
        if oran_gun:
            s["tahmin"] = yuvarla(net * oran_gun[0] / oran_gun[1])
            s["tahmin_prim"] = prim_hesapla(s["tahmin"], a["kademe"], a["yontem"])[0]
            if sonraki and s["tahmin"] >= sonraki[0]:
                uy.append(uyari("Bilgi", "Sonraki kademeye ulaşılabilir", a["tedarikci"], f"Bu hızla dönem sonu cirosu {s['tahmin']:,.2f}; {sonraki[0]:,.2f} kademesine (%{sonraki[1]}) ulaşılır."))
        if k not in ciro:
            uy.append(uyari("Bilgi", "Alım yok", a["tedarikci"], "Anlaşması olan tedarikçiden dönemde alım yapılmamış."))
        elif net < 0:
            uy.append(uyari("Orta", "İade alımdan fazla", a["tedarikci"], f"Net ciro {net:,.2f}; prim hesaplanmadı."))
        elif sonraki and s["kalan"] <= sonraki[0] * Decimal("0.1"):
            uy.append(uyari("Bilgi", "Sonraki kademeye yakın", a["tedarikci"], f"{s['kalan']:,.2f} daha alımla %{sonraki[1]} kademesine geçilir; ek prim {s['ek_prim']:,.2f}."))
        sonuc.append(s)
    for k in sorted(set(ciro) - set(anlasmalar)):
        uy.append(uyari("Bilgi", "Anlaşma yok", adlar[k], f"Dönemde {ciro[k]:,.2f} alım var, ciro primi anlaşması yok."))
    aylar = sorted({ay for _, ay in aylik})
    return {"tedarikciler": sonuc, "aylar": aylar, "aylik": aylik, "anlasmalar": anlasmalar, "toplam_prim": sum((s["prim"] for s in sonuc), SIFIR), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Prim Hesabı"
    _baslik(ws, ["Tedarikçi", "Yöntem", "Alımlar", "İadeler", "Net Ciro", "Ulaşılan Kademe", "Prim %", "Hak Edilen Prim", "Sonraki Kademe", "Sonraki Prim %", "Kalan Ciro", "Ulaşılırsa Ek Prim",
                "Dönem Sonu Tahmini Ciro", "Tahmini Prim"], (24, 10, 15, 13, 15, 14, 8, 14, 14, 9, 14, 13, 15, 13), "B2")
    for s in r["tedarikciler"]:
        _satir(ws, [s["tedarikci"], s["yontem"], s["brut"], s["iade"] or None, s["ciro"], s["alt"], s["oran"], s["prim"], s["sonraki_alt"], s["sonraki_oran"], s["kalan"], s["ek_prim"], s["tahmin"], s["tahmin_prim"]],
               para_s=(3, 4, 5, 6, 8, 9, 11, 12, 13, 14), dolgu={8: YESIL if s["prim"] else None})
    k = _satir(ws, ["TOPLAM", None, None, None, sum((s["ciro"] for s in r["tedarikciler"]), SIFIR), None, None, r["toplam_prim"]], para_s=(5, 8))
    for h in ws[k]:
        h.font = Font(bold=True)
    ws.append([])
    ws.append(["Tutarlar KDV hariçtir. Prim faturası, KDV ve muhasebe kaydı ayrıca yapılır; ciro tanımını (iade, iskonto, fiyat farkı dahil mi) anlaşmadan doğrulayın."])
    ws = wb.create_sheet("Kademeler")
    _baslik(ws, ["Tedarikçi", "Yöntem", "Alt Sınır", "Prim %"], (24, 10, 15, 8))
    for _, a in sorted(r["anlasmalar"].items()):
        for alt, oran in a["kademe"]:
            _satir(ws, [a["tedarikci"], a["yontem"], alt, oran], para_s=(3, 4))
    if r["aylar"]:
        ws = wb.create_sheet("Aylık Ciro")
        _baslik(ws, ["Tedarikçi"] + r["aylar"], (24,) + (13,) * len(r["aylar"]), "B2")
        adlar = {katla(s["tedarikci"]): s["tedarikci"] for s in r["tedarikciler"]}
        for k in sorted({k for k, _ in r["aylik"]}):
            _satir(ws, [adlar.get(k, k)] + [r["aylik"].get((k, ay)) for ay in r["aylar"]], para_s=range(2, len(r["aylar"]) + 2))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(anlasma_yol: Path, alim_yol: Path, cikti: Path, baslangic: date | None = None, bitis: date | None = None, gun: date | None = None) -> dict:
    anlasmalar, uy1 = oku_anlasma(anlasma_yol)
    alimlar, uy2 = oku_alim(alim_yol)
    r = hesapla(anlasmalar, alimlar, baslangic, bitis, gun)
    r["uyarilar"] = sirala(uy1 + uy2 + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Tedarikçi anlaşmalarındaki ciro primi kademelerine göre hak edilen primleri hesaplar.")
    p.add_argument("--anlasmalar", type=Path, help="Tedarikçi, Alt Sınır, Prim %%, Yöntem (Tüm ciro / Dilimli)")
    p.add_argument("--alimlar", type=Path, help="Tedarikçi, Tarih, Tutar (KDV hariç; iade eksi ya da Tür = İade)")
    p.add_argument("--baslangic", help="Anlaşma dönemi başı GG.AA.YYYY")
    p.add_argument("--bitis", help="Anlaşma dönemi sonu GG.AA.YYYY")
    p.add_argument("--tarih", help="Verinin kesildiği gün; dönem sonu tahmini için GG.AA.YYYY")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "tedarikci_ciro_primi.xlsx")
    a = p.parse_args(argv)
    ornek = a.anlasmalar is None and a.alimlar is None
    ay, ly = a.anlasmalar or ORNEK / "ciro_primi_anlasmalari.csv", a.alimlar or ORNEK / "tedarikci_alimlari.csv"
    if ornek and not (a.baslangic or a.bitis or a.tarih):
        a.baslangic, a.bitis, a.tarih = "01.01.2026", "31.12.2026", "30.09.2026"
    tarihler = [tarih(x) if x else None for x in (a.baslangic, a.bitis, a.tarih)]
    if any(x and t is None for x, t in zip((a.baslangic, a.bitis, a.tarih), tarihler)) or (tarihler[0] and tarihler[1] and tarihler[0] > tarihler[1]):
        print("[X] Tarihler GG.AA.YYYY olmalı; --baslangic --bitis'ten sonra olamaz")
        return 2
    for y in (ay, ly):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(ay, ly, a.cikti, *tarihler)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(r['tedarikciler'])} tedarikçi · hak edilen prim toplamı {r['toplam_prim']:,.2f}")
    for s in r["tedarikciler"]:
        print(f"     {s['tedarikci']}: net ciro {s['ciro']:,.2f} · prim {s['prim']:,.2f}" + (f" (%{s['oran']})" if s["oran"] is not None else "")
              + (f" · sonraki kademeye {s['kalan']:,.2f}" if s["kalan"] is not None else ""))
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
