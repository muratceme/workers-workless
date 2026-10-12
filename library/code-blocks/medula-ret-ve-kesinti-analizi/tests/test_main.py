import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def k(hata, tutar, ay=7, islem="I-1", brans="B"):
    return {"donem": f"2026-{ay:02d}", "donem_sira": (2026, ay), "brans": brans, "islem_kodu": islem, "islem": "", "hata_kodu": hata, "hata": f"{hata} açıklama", "tutar": D(tutar)}


class AnalizTesti(unittest.TestCase):
    def test_donem_cozme(self):
        self.assertEqual([main.donem_coz(x) for x in ("2026-08", "08.2026", "2026/8", "Ağustos 2026", "EYLÜL 2026", "01.08.2026", date(2026, 8, 15), "2026-13", "3. çeyrek", "")],
                         [(2026, 8), (2026, 8), (2026, 8), (2026, 8), (2026, 9), (2026, 8), (2026, 8), None, None, None])

    def test_hata_bazinda_ve_tekrarlayan(self):
        l = [k("A", 100, 7), k("A", 100, 8), k("A", 100, 9), k("B", 500, 7), k("B", 500, 7), k("B", 500, 7), k("C", 50, 7), k("C", 50, 8)]
        r = main.analiz(l)
        self.assertEqual([(h["hata_kodu"], h["adet"], h["tutar"], h["pay"], h["birikimli"], h["donem_sayisi"], h["tekrarlayan"]) for h in r["hatalar"]],
                         [("B", 3, D(1500), D("78.95"), D("78.95"), 1, False), ("A", 3, D(300), D("15.79"), D("94.74"), 3, True), ("C", 2, D(100), D("5.26"), D("100.00"), 2, False)])
        self.assertEqual(r["toplam"], {"adet": 8, "tutar": D(1900), "tekrarlayan": 1, "tekrarlayan_tutar": D(300)})
        self.assertEqual([h["tekrarlayan"] for h in main.analiz(l, min_donem=1, min_adet=2)["hatalar"]], [True, True, True])
        self.assertEqual(r["hatalar"][0]["hata"], "B açıklama")
        with self.assertRaises(ValueError):
            main.analiz([])

    def test_cift_brans_ve_donem(self):
        l = [k("A", 100, 7, "I-1"), k("A", 100, 8, "I-1"), k("A", 100, 9, "I-1"), k("A", 300, 9, "I-2"), k("B", 500, 8, "I-3", "C"), k("B", 0, 9, "I-3", "C")]
        r = main.analiz(l)
        self.assertEqual([(c["hata_kodu"], c["islem"], c["adet"], c["tutar"], c["donem_sayisi"]) for c in r["ciftler"]], [("A", "I-1", 3, D(300), 3)])
        self.assertEqual([(b["brans"], b["adet"], b["tutar"], b["pay"], b["en_sik"]) for b in r["branslar"]], [("B", 4, D(600), D("54.55"), "A"), ("C", 2, D(500), D("45.45"), "B")])
        self.assertEqual([(d["donem"], d["adet"], d["tutar"], d["degisim"]) for d in r["donemler"]], [("2026-07", 1, D(100), None), ("2026-08", 2, D(600), D("500.00")), ("2026-09", 3, D(400), D("-33.33"))])
        self.assertEqual([(i["islem"], i["tutar"]) for i in r["islemler"]], [("I-3", D(500)), ("I-1", D(300)), ("I-2", D(300))])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "m.xlsx"
            r = main.calistir(ORNEK / "ret_kesinti_kayitlari.csv", cikti)
            self.assertEqual([(h["hata_kodu"], h["adet"], h["tutar"], h["pay"], h["birikimli"], h["donem_sayisi"], h["tekrarlayan"]) for h in r["hatalar"]],
                             [("H-02", 3, D(4500), D("66.18"), D("66.18"), 3, True), ("H-01", 5, D(1100), D("16.18"), D("82.35"), 3, True), ("H-03", 1, D(800), D("11.76"), D("94.12"), 1, False),
                              ("H-04", 1, D(400), D("5.88"), D("100.00"), 1, False)])
            self.assertEqual([(b["brans"], b["adet"], b["tutar"], b["pay"], b["en_sik"]) for b in r["branslar"]],
                             [("Ortopedi", 4, D(5300), D("77.94"), "H-02"), ("Dahiliye", 5, D(1100), D("16.18"), "H-01"), ("Kardiyoloji", 1, D(400), D("5.88"), "H-04")])
            self.assertEqual([(d["donem"], d["adet"], d["tutar"], d["degisim"]) for d in r["donemler"]], [("2026-07", 3, D(1900), None), ("2026-08", 3, D(2500), D("31.58")), ("2026-09", 4, D(2400), D("-4.00"))])
            self.assertEqual([(c["hata_kodu"], c["islem"], c["adet"], c["tutar"]) for c in r["ciftler"]], [("H-01", "I-100 · Muayene", 4, D(800)), ("H-02", "I-300 · Görüntüleme", 3, D(4500))])
            self.assertEqual((r["toplam"], r["uyarilar"]), ({"adet": 10, "tutar": D(6800), "tekrarlayan": 2, "tekrarlayan_tutar": D(5600)}, []))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Hata Kodu Bazında", "Tekrarlayan Hata-İşlem", "Branş Bazında", "İşlem Bazında", "Dönem Eğilimi", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "k.csv"
            y.write_text("Dönem;Ret Kodu;Açıklama;Tutar\nEylül 2026;A;x;1.250,50\n3. çeyrek;A;x;10\n3. çeyrek;B;y;20\n09.2026;C;z;çok\n;D;;5\n09.2026;;;5\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual([(x["donem"], x["hata_kodu"], x["tutar"], x["brans"]) for x in l],
                             [("2026-09", "A", D("1250.50"), "Belirtilmemiş"), ("3. çeyrek", "A", D(10), "Belirtilmemiş"), ("3. çeyrek", "B", D(20), "Belirtilmemiş"), ("Belirtilmemiş", "D", D(5), "Belirtilmemiş")])
            self.assertEqual(sorted(u["tur"] for u in uy), ["Dönem biçimi tanınmadı", "Tutar okunamadı"])
            self.assertEqual([d["donem"] for d in main.analiz(l)["donemler"]][0], "2026-09")
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--kayitlar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--min-donem", "1", "--min-adet", "2"]), 0)
            self.assertEqual(main.main(arg + ["--min-adet", "0"]), 2)
            self.assertEqual(main.main(["--kayitlar", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Hata Kodu;Tutar\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            y.write_text("Dönem;Hata Kodu;Tutar\n2026-09;A;5\n", encoding="utf-8")
            self.assertEqual([u["tur"] for u in main.calistir(y, Path(c))["uyarilar"]], ["Dönem sayısı az"])
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
