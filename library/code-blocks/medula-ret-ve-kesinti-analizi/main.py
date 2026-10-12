"""
Medula Ret ve Kesinti Analizi — Workers / Workless kod bloğu
Anlaşmalı Kurumlar ve Faturalama › SGK Faturalama Uzmanı

Medula ret/kesinti kayıtlarını hata kodu, branş ve işlem bazında analiz edip tekrarlayan hataları raporlar.
  Hata kodu bazında  = adet, tutar, tutar payı %, birikimli pay % (tutara göre büyükten küçüğe: Pareto), görüldüğü dönem sayısı
  Tekrarlayan hata   = hata kodu en az --min-donem (2) farklı dönemde ve toplam en az --min-adet (3) kez görülmüş
  Tekrarlayan çift   = aynı hata kodu + aynı işlem en az --min-adet kez görülmüş (nereden düzeltileceğini gösterir)
  Dönem eğilimi      = dönem bazında adet ve tutar; önceki döneme göre tutar değişimi %
Medula'ya bağlanmaz; sizin dışa aktardığınız ret/kesinti listesini sayar. Hata kodlarının anlamını ve SUT kurallarını bilmez; açıklamalar dosyanızdan gelir.
Hasta bilgisi gerekmez: T.C. kimlik no, ad ve takip numarası sütunlarını dosyaya koymayın.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"donem": ("donem", "fatura donemi", "ay"), "brans": ("brans", "bolum", "klinik"), "islem_kodu": ("islem kodu", "sut kodu", "hizmet kodu", "kod"), "islem": ("islem adi", "islem", "hizmet adi", "hizmet"),
         "hata_kodu": ("hata kodu", "ret kodu", "kesinti kodu", "red kodu"), "hata": ("hata aciklamasi", "hata", "aciklama", "ret nedeni", "kesinti nedeni"), "tutar": ("tutar", "kesinti tutari", "ret tutari")}
AYLAR = ("ocak", "subat", "mart", "nisan", "mayis", "haziran", "temmuz", "agustos", "eylul", "ekim", "kasim", "aralik")
Z = Decimal(0)


def donem_coz(x) -> tuple[int, int] | None:
    """2026-08 · 08.2026 · 2026/8 · Ağustos 2026 · tarih → (yıl, ay)"""
    if isinstance(x, (date, datetime)):
        return x.year, x.month
    k = katla(x)
    m = re.fullmatch(r"(\d{4}) (\d{1,2})", k) or re.fullmatch(r"(\d{4}) (\d{1,2}) \d{1,2}.*", k)
    if m and 1 <= int(m[2]) <= 12:
        return int(m[1]), int(m[2])
    m = re.fullmatch(r"(?:\d{1,2} )?(\d{1,2}) (\d{4})", k)
    if m and 1 <= int(m[1]) <= 12:
        return int(m[2]), int(m[1])
    m = re.fullmatch(r"([a-z]+) (\d{4})", k)
    if m and m[1] in AYLAR:
        return int(m[2]), AYLAR.index(m[1]) + 1
    return None


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, bilinmeyen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("hata_kodu", "tutar")):
        h = metin(r["hata_kodu"])
        if not h:
            continue
        kim = f"{h} (satır {r['_satir']})"
        tutar = para(r["tutar"])
        if tutar is None or tutar < 0:
            uy.append(uyari("Yüksek", "Tutar okunamadı", kim, f"'{metin(r['tutar'])}'. Satır alınmadı."))
            continue
        d = donem_coz(r["donem"])
        ham = metin(r["donem"])
        if d is None and ham and ham not in bilinmeyen:
            bilinmeyen.add(ham)
            uy.append(uyari("Orta", "Dönem biçimi tanınmadı", ham, "Dönem metin olarak alındı ve sıralamada sona kondu. Tanınan biçimler: 2026-08, 08.2026, Ağustos 2026."))
        l.append({"donem": f"{d[0]}-{d[1]:02d}" if d else (ham or "Belirtilmemiş"), "donem_sira": d or (9999, len(bilinmeyen)), "brans": metin(r["brans"]) or "Belirtilmemiş",
                  "islem_kodu": metin(r["islem_kodu"]), "islem": metin(r["islem"]), "hata_kodu": h, "hata": metin(r["hata"]), "tutar": tutar})
    return l, uy


def analiz(kayit: list[dict], min_donem: int = 2, min_adet: int = 3) -> dict:
    if not kayit:
        raise ValueError("Analiz edilecek ret/kesinti kaydı yok.")
    toplam = sum((k["tutar"] for k in kayit), Z)

    def grupla(anahtar, ad) -> list[dict]:
        g = {}
        for k in kayit:
            x = g.setdefault(anahtar(k), ad(k) | {"adet": 0, "tutar": Z, "donemler": set(), "hatalar": defaultdict(int)})
            x["adet"] += 1
            x["tutar"] += k["tutar"]
            x["donemler"].add(k["donem"])
            x["hatalar"][k["hata_kodu"]] += 1
        l, bir = sorted(g.values(), key=lambda x: (-x["tutar"], -x["adet"])), Z
        for x in l:
            bir += x["tutar"]
            x["pay"] = yuvarla(x["tutar"] / toplam * 100) if toplam else Z
            x["birikimli"] = yuvarla(bir / toplam * 100) if toplam else Z
            x["donem_sayisi"] = len(x["donemler"])
            x["en_sik"] = max(x["hatalar"], key=lambda h: (x["hatalar"][h], h))
        return l

    islem_ad = lambda k: " · ".join(x for x in (k["islem_kodu"], k["islem"]) if x) or "Belirtilmemiş"   # noqa: E731
    hatalar = grupla(lambda k: katla(k["hata_kodu"]), lambda k: {"hata_kodu": k["hata_kodu"], "hata": k["hata"]})
    for h in hatalar:
        h["tekrarlayan"] = h["donem_sayisi"] >= min_donem and h["adet"] >= min_adet
    ciftler = [c for c in grupla(lambda k: (katla(k["hata_kodu"]), katla(k["islem_kodu"]) or katla(k["islem"])), lambda k: {"hata_kodu": k["hata_kodu"], "hata": k["hata"], "islem": islem_ad(k), "brans": k["brans"]})
               if c["adet"] >= min_adet]
    donemler, sira = grupla(lambda k: k["donem"], lambda k: {"donem": k["donem"]}), {k["donem"]: k["donem_sira"] for k in kayit}
    donemler.sort(key=lambda d: sira[d["donem"]])
    for onceki, d in zip([None, *donemler], donemler):
        d["degisim"] = yuvarla((d["tutar"] - onceki["tutar"]) / onceki["tutar"] * 100) if onceki and onceki["tutar"] else None
    return {"hatalar": hatalar, "branslar": grupla(lambda k: katla(k["brans"]), lambda k: {"brans": k["brans"]}),
            "islemler": grupla(lambda k: katla(k["islem_kodu"]) or katla(k["islem"]), lambda k: {"islem": islem_ad(k)}), "ciftler": sorted(ciftler, key=lambda c: (-c["adet"], -c["tutar"])),
            "donemler": donemler, "toplam": {"adet": len(kayit), "tutar": toplam, "tekrarlayan": sum(h["tekrarlayan"] for h in hatalar),
                                             "tekrarlayan_tutar": sum((h["tutar"] for h in hatalar if h["tekrarlayan"]), Z)}, "ayar": (min_donem, min_adet)}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    md, ma = r["ayar"]
    t = r["toplam"]
    ws = wb.active
    ws.title = "Hata Kodu Bazında"
    _baslik(ws, ["Hata Kodu", "Hata Açıklaması", "Adet", "Tutar", "Tutar Payı %", "Birikimli %", "Görüldüğü Dönem", "Tekrarlayan", "Düzeltici Aksiyon (siz doldurun)"], (14, 50, 8, 15, 10, 10, 10, 12, 36))
    for h in r["hatalar"]:
        _satir(ws, [h["hata_kodu"], h["hata"], h["adet"], h["tutar"], h["pay"], h["birikimli"], h["donem_sayisi"], "Evet" if h["tekrarlayan"] else "", ""], para_s=(4, 5, 6),
               dolgu={8: KIRMIZI if h["tekrarlayan"] else None}, sar=(2,))
    _satir(ws, ["Toplam", "", t["adet"], t["tutar"], 100], para_s=(4, 5))
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.append([])
    ws.append([f"Tekrarlayan: en az {md} farklı dönemde ve toplam en az {ma} kez görülen hata kodu. {t['tekrarlayan']} tekrarlayan hata kodu toplam tutarın {t['tekrarlayan_tutar']:,.2f} kısmını oluşturuyor. "
               "Hata açıklamaları dosyanızdan alınmıştır; kod SUT kurallarını değerlendirmez."])
    ws = wb.create_sheet("Tekrarlayan Hata-İşlem")
    _baslik(ws, ["Hata Kodu", "Hata Açıklaması", "İşlem", "Branş", "Adet", "Tutar", "Görüldüğü Dönem"], (14, 44, 36, 20, 8, 15, 10))
    for c in r["ciftler"]:
        _satir(ws, [c["hata_kodu"], c["hata"], c["islem"], c["brans"], c["adet"], c["tutar"], c["donem_sayisi"]], para_s=(6,), sar=(2,))
    for ad, alan, bas, liste in (("Branş Bazında", "brans", "Branş", r["branslar"]), ("İşlem Bazında", "islem", "İşlem", r["islemler"])):
        ws = wb.create_sheet(ad)
        _baslik(ws, [bas, "Adet", "Tutar", "Tutar Payı %", "Birikimli %", "En Sık Hata Kodu"], (40, 8, 15, 10, 10, 16))
        for x in liste:
            _satir(ws, [x[alan], x["adet"], x["tutar"], x["pay"], x["birikimli"], x["en_sik"]], para_s=(3, 4, 5))
    ws = wb.create_sheet("Dönem Eğilimi")
    _baslik(ws, ["Dönem", "Adet", "Tutar", "Önceki Döneme Göre Tutar Değişimi %", "En Sık Hata Kodu"], (14, 8, 15, 18, 16))
    for d in r["donemler"]:
        _satir(ws, [d["donem"], d["adet"], d["tutar"], d["degisim"], d["en_sik"]], para_s=(3, 4), dolgu={4: (KIRMIZI if d["degisim"] > 0 else YESIL) if d["degisim"] else None})
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, min_donem: int = 2, min_adet: int = 3) -> dict:
    kayit, uy = oku(yol)
    r = analiz(kayit, min_donem, min_adet)
    if len(r["donemler"]) < min_donem:
        uy.append(uyari("Bilgi", "Dönem sayısı az", yol.name, f"Veride {len(r['donemler'])} dönem var; tekrarlayan hata için en az {min_donem} dönem gerekir. Birkaç dönemin kayıtlarını birlikte verin."))
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Ret/kesinti kayıtlarını hata kodu, branş ve işlem bazında analiz eder.")
    p.add_argument("--kayitlar", type=Path, default=ORNEK / "ret_kesinti_kayitlari.csv", help="Dönem, Branş, İşlem Kodu, İşlem Adı, Hata Kodu, Hata Açıklaması, Tutar")
    p.add_argument("--min-donem", default="2", help="Tekrarlayan sayılması için hata kodunun görülmesi gereken en az dönem sayısı (varsayılan 2)")
    p.add_argument("--min-adet", default="3", help="Tekrarlayan sayılması için en az kayıt sayısı (varsayılan 3)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "ret_kesinti_analizi.xlsx")
    a = p.parse_args(argv)
    md, ma = para(a.min_donem), para(a.min_adet)
    if any(x is None or x < 1 or x != int(x) for x in (md, ma)):
        print("[X] --min-donem ve --min-adet 1 veya daha büyük tam sayı olmalı")
        return 2
    if not a.kayitlar.exists():
        print(f"[X] Dosya bulunamadı: {a.kayitlar}")
        return 1
    try:
        r = calistir(a.kayitlar, a.cikti, int(md), int(ma))
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {t['adet']} kayıt · {t['tutar']:,.2f} · {len(r['hatalar'])} hata kodu · {len(r['donemler'])} dönem")
    for h in r["hatalar"]:
        if h["tekrarlayan"]:
            print(f"[!] Tekrarlayan {h['hata_kodu']}: {h['adet']} kayıt, {h['tutar']:,.2f} (%{h['pay']}), {h['donem_sayisi']} dönem" + (f" · {h['hata']}" if h["hata"] else ""))
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
