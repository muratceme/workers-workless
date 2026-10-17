"""
Parti Verim ve Fire Raporu — Gerisi Bende kod bloğu
Üretim › Üretim Şefi

Parti bazında hammadde girişi, mamul çıkışı ve fireden verimi hesaplar; hedef dışı partileri, ürün ve hat özetini çıkarır.
  Verim %            = mamul çıkışı ÷ hammadde girişi × 100
  Fire %             = fire ÷ hammadde girişi × 100;   yeniden işleme % = yeniden işlemeye ayrılan ÷ giriş × 100
  Kayıt dışı fark    = giriş − mamul − fire − yeniden işleme                 (tartılmayan kayıp: nem, buharlaşma, dökülme, tartım farkı)
  Hedeften sapma     = verim − hedef verim (puan);  sapma < −tolerans → "Hedefin altında",  > +tolerans → "Hedefin üzerinde"
  Kayıp mamul        = (hedef − verim) ÷ 100 × giriş                         (yalnız hedefin altındaki partilerde)
  Ürün / hat verimi  = Σ mamul ÷ Σ giriş × 100                               (partilerin basit ortalaması değil, ağırlıklı)
Giriş ve çıkış aynı birimde (ör. kg) olmalıdır. Pişirme, kurutma gibi nem kaybı olan proseslerde düşük verim normaldir; hedefi buna göre siz belirlersiniz.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"parti": ("parti no", "parti", "uretim partisi", "lot no"), "urun": ("urun", "urun adi", "mamul adi"), "tarih": ("tarih", "uretim tarihi"), "hat": ("hat", "uretim hatti", "hat vardiya", "vardiya"),
         "giris": ("hammadde girisi", "hammadde girisi kg", "giris", "giris kg", "girdi", "hammadde"), "cikis": ("mamul cikisi", "mamul cikisi kg", "cikis", "cikis kg", "mamul", "uretim miktari"),
         "fire": ("fire", "fire kg", "iskarta"), "rework": ("yeniden isleme", "yeniden isleme kg", "rework", "geri kazanim")}
HEDEF = {"urun": SUTUN["urun"], "hedef": ("hedef verim", "hedef", "hedef verim yuzde")}
Z = Decimal(0)
DURUM_RENK = {"Hedefin altında": KIRMIZI, "Hedefin üzerinde": SARI, "Hedefte": YESIL}


def olcu(x) -> Decimal | None:
    """Miktar: 12,5 = 12.5; 1.250,5 = bin iki yüz elli buçuk. Tek başına nokta ondalık ayırıcıdır."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    else:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def oku(yol: Path, hedef_yol: Path | None = None) -> tuple[list[dict], dict, list[dict]]:
    uy, l, hedef = [], [], {}
    for r in kayitlar(yol, SUTUN, ("parti", "giris", "cikis")):
        p = metin(r["parti"])
        if not p:
            continue
        g, c, f, rw = olcu(r["giris"]), olcu(r["cikis"]), olcu(r["fire"]) if metin(r["fire"]) else Z, olcu(r["rework"]) if metin(r["rework"]) else Z
        if g is None or g <= 0 or c is None or c < 0 or f is None or f < 0 or rw is None or rw < 0:
            uy.append(uyari("Yüksek", "Parti okunamadı", f"{p} (satır {r['_satir']})", f"Giriş '{metin(r['giris'])}', çıkış '{metin(r['cikis'])}', fire '{metin(r['fire'])}', yeniden işleme '{metin(r['rework'])}'. "
                            "Giriş pozitif, diğerleri sıfır veya pozitif olmalı. Satır alınmadı."))
            continue
        l.append({"parti": p, "urun": metin(r["urun"]) or "Belirtilmemiş", "tarih": tarih(r["tarih"]), "hat": metin(r["hat"]) or "Belirtilmemiş", "giris": g, "cikis": c, "fire": f, "rework": rw})
    if hedef_yol:
        for r in kayitlar(hedef_yol, HEDEF, ("urun", "hedef")):
            u = metin(r["urun"])
            if not u:
                continue
            h = olcu(r["hedef"])
            if h is None or not 0 < h <= 100:
                uy.append(uyari("Orta", "Hedef okunamadı", u, f"Hedef verim '{metin(r['hedef'])}' 0–100 arasında olmalı. Bu ürün hedefsiz değerlendirildi."))
                continue
            hedef[katla(u)] = h
    return l, hedef, uy


def yuzde(a: Decimal, b: Decimal) -> Decimal | None:
    return yuvarla(a / b * 100) if b else None


def analiz(partiler: list[dict], hedefler: dict | None = None, hedef_genel: Decimal | None = None, tolerans: Decimal = Decimal(2), fark_esik: Decimal = Decimal(5)) -> dict:
    if not partiler:
        raise ValueError("Değerlendirilecek parti yok.")
    hedefler, uy, liste = hedefler or {}, [], []
    grup = {"urun": {}, "hat": {}}
    for p in partiler:
        g, c = p["giris"], p["cikis"]
        verim, fark = c / g * 100, g - c - p["fire"] - p["rework"]
        hedef = hedefler.get(katla(p["urun"]), hedef_genel)
        sapma = verim - hedef if hedef is not None else None
        durum = "Hedef yok" if hedef is None else "Hedefin altında" if sapma < -tolerans else "Hedefin üzerinde" if sapma > tolerans else "Hedefte"
        kayip = yuvarla((hedef - verim) / 100 * g) if durum == "Hedefin altında" else Z
        x = p | {"verim": yuvarla(verim), "fire_y": yuzde(p["fire"], g), "rework_y": yuzde(p["rework"], g), "fark": fark, "fark_y": yuzde(fark, g), "hedef": hedef, "sapma": yuvarla(sapma) if sapma is not None else None,
                 "durum": durum, "kayip": kayip}
        liste.append(x)
        if verim > 100:
            uy.append(uyari("Yüksek", "Verim %100'ün üzerinde", p["parti"], f"Giriş {g:f}, mamul {c:f} (verim %{x['verim']}). Girişe su veya katkı eklenip yazılmamış ya da tartım, parti eşleştirme hatası olabilir."))
        elif fark < 0:
            uy.append(uyari("Orta", "Çıkış toplamı girişten fazla", p["parti"], f"Mamul + fire + yeniden işleme, girişten {-fark:f} fazla. Kayıtları kontrol edin."))
        elif x["fark_y"] > fark_esik:
            uy.append(uyari("Orta", "Kayıt dışı fark yüksek", p["parti"], f"Giriş − mamul − fire − yeniden işleme = {fark:f} (%{x['fark_y']}). Tartılmayan fire, nem kaybı veya eksik kayıt olabilir."))
        if durum == "Hedefin altında":
            uy.append(uyari("Orta", "Hedefin altında", p["parti"], f"{p['urun']}: verim %{x['verim']}, hedef %{hedef.normalize():f} (sapma {x['sapma']} puan, yaklaşık {kayip:f} birim mamul kaybı)."))
        for tur, anahtar in (("urun", p["urun"]), ("hat", p["hat"])):
            o = grup[tur].setdefault(katla(anahtar), {"ad": anahtar, "parti": 0, "giris": Z, "cikis": Z, "fire": Z, "rework": Z, "altinda": 0, "kayip": Z, "verimler": [], "hedef": hedef if tur == "urun" else None})
            o["parti"] += 1
            o["altinda"] += durum == "Hedefin altında"
            o["kayip"] += kayip
            o["verimler"].append(x["verim"])
            for a in ("giris", "cikis", "fire", "rework"):
                o[a] += p[a]

    def kapat(o: dict) -> dict:
        v = o.pop("verimler")
        return o | {"verim": yuzde(o["cikis"], o["giris"]), "fire_y": yuzde(o["fire"], o["giris"]), "en_dusuk": min(v), "en_yuksek": max(v)}

    genel = {"ad": "Genel", "parti": len(liste), "altinda": sum(1 for x in liste if x["durum"] == "Hedefin altında"), "kayip": sum((x["kayip"] for x in liste), Z), "verimler": [x["verim"] for x in liste], "hedef": None}
    for a in ("giris", "cikis", "fire", "rework"):
        genel[a] = sum((x[a] for x in liste), Z)
    say = defaultdict(int)
    for x in liste:
        say[x["durum"]] += 1
    return {"partiler": liste, "urunler": sorted((kapat(o) for o in grup["urun"].values()), key=lambda o: o["ad"]), "hatlar": sorted((kapat(o) for o in grup["hat"].values()), key=lambda o: o["ad"]),
            "genel": kapat(genel), "say": dict(say), "ayar": (tolerans, fark_esik), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Partiler"
    _baslik(ws, ["Parti No", "Ürün", "Tarih", "Hat", "Giriş", "Mamul", "Fire", "Yeniden İşleme", "Verim %", "Fire %", "Kayıt Dışı Fark", "Fark %", "Hedef %", "Sapma (puan)", "Durum", "Kayıp Mamul", "Neden (siz doldurun)"],
            (14, 26, 12, 14, 12, 12, 10, 11, 9, 9, 11, 8, 8, 9, 18, 11, 30))
    for x in r["partiler"]:
        _satir(ws, [x["parti"], x["urun"], x["tarih"], x["hat"], x["giris"], x["cikis"], x["fire"], x["rework"] or None, x["verim"], x["fire_y"], x["fark"], x["fark_y"], x["hedef"], x["sapma"], x["durum"], x["kayip"] or None, ""],
               para_s=(5, 6, 7, 8, 9, 10, 11, 12, 14, 16), tarih_s=(3,), dolgu={15: DURUM_RENK.get(x["durum"]), 9: KIRMIZI if x["verim"] > 100 else None})
    tol, fe = r["ayar"]
    ws.append([])
    ws.append([f"Verim = mamul ÷ giriş × 100. Sapma ±{tol.normalize():f} puanı aşarsa hedef dışı. Kayıt dışı fark = giriş − mamul − fire − yeniden işleme; %{fe.normalize():f} üzeri işaretlenir."])
    for ad, satirlar, baslik in (("Ürün Özeti", r["urunler"] + [r["genel"]], "Ürün"), ("Hat Özeti", r["hatlar"], "Hat")):
        ws = wb.create_sheet(ad)
        _baslik(ws, [baslik, "Parti Sayısı", "Giriş", "Mamul", "Fire", "Verim %", "Fire %", "En Düşük Verim %", "En Yüksek Verim %", "Hedef %", "Hedef Altı Parti", "Kayıp Mamul"], (26, 9, 13, 13, 11, 9, 9, 11, 11, 9, 10, 12))
        for o in satirlar:
            n = _satir(ws, [o["ad"], o["parti"], o["giris"], o["cikis"], o["fire"], o["verim"], o["fire_y"], o["en_dusuk"], o["en_yuksek"], o["hedef"], o["altinda"] or None, o["kayip"] or None], para_s=(3, 4, 5, 6, 7, 8, 9, 12),
                       dolgu={11: KIRMIZI if o["altinda"] else None})
            if o["ad"] == "Genel" and o is r["genel"]:
                for c in ws[n]:
                    c.font = Font(bold=True)
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, hedef_yol: Path | None = None, **ayar) -> dict:
    partiler, hedefler, uy = oku(yol, hedef_yol)
    r = analiz(partiler, hedefler, **ayar)
    uy += r["uyarilar"]
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Parti bazında verim ve fireyi hesaplar, hedef dışı partileri listeler.")
    p.add_argument("--partiler", type=Path, default=ORNEK / "partiler.csv", help="Parti No, Ürün, Tarih, Hat, Hammadde Girişi, Mamul Çıkışı, Fire, Yeniden İşleme")
    p.add_argument("--hedefler", type=Path, help="İsteğe bağlı: Ürün, Hedef Verim %%")
    p.add_argument("--hedef", help="Hedef dosyasında olmayan ürünler için genel hedef verim %%")
    p.add_argument("--tolerans", default="2", help="Hedeften sapma toleransı, puan (varsayılan 2)")
    p.add_argument("--fark", default="5", help="Kayıt dışı fark bu %%'yi aşarsa uyarı (varsayılan 5)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "parti_verim_raporu.xlsx")
    a = p.parse_args(argv)
    hedef_yol = a.hedefler or (ORNEK / "hedefler.csv" if a.partiler == ORNEK / "partiler.csv" else None)
    h, t, f = (olcu(a.hedef) if a.hedef else None), olcu(a.tolerans), olcu(a.fark)
    if (a.hedef and (h is None or not 0 < h <= 100)) or t is None or t < 0 or f is None or f < 0:
        print("[X] --hedef 0–100 arasında; --tolerans ve --fark sıfır veya pozitif sayı olmalı")
        return 2
    for y in (a.partiler, hedef_yol):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.partiler, a.cikti, hedef_yol, hedef_genel=h, tolerans=t, fark_esik=f)
    except ValueError as hata:
        print(f"[X] {hata}")
        return 1
    g = r["genel"]
    print(f"[OK] {g['parti']} parti · giriş {g['giris']:,.2f} · mamul {g['cikis']:,.2f} · verim %{g['verim']} · fire %{g['fire_y']} · {g['altinda']} parti hedefin altında")
    for x in r["partiler"]:
        if x["durum"] == "Hedefin altında":
            print(f"[!] {x['parti']} ({x['urun']}): verim %{x['verim']}, hedef %{x['hedef'].normalize():f}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
