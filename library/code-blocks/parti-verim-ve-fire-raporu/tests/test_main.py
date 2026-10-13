import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def p(parti, giris, cikis, fire=0, rework=0, urun="U", hat="H"):
    return {"parti": parti, "urun": urun, "tarih": None, "hat": hat, "giris": D(giris), "cikis": D(cikis), "fire": D(fire), "rework": D(rework)}


class AnalizTesti(unittest.TestCase):
    def test_parti_hesabi(self):
        x = main.analiz([p("P", 1200, 1110, 60, 20)], {"u": D(90)})["partiler"][0]
        self.assertEqual((x["verim"], x["fire_y"], x["rework_y"], x["fark"], x["fark_y"], x["hedef"], x["sapma"], x["durum"], x["kayip"]),
                         (D("92.50"), D("5.00"), D("1.67"), D(10), D("0.83"), D(90), D("2.50"), "Hedefin üzerinde", D(0)))
        d = lambda c, **kw: main.analiz([p("P", 1000, c)], {"u": D(90)}, **kw)["partiler"][0]   # noqa: E731
        self.assertEqual([(d(c)["durum"], d(c)["kayip"]) for c in (879, 880, 900, 920, 921)],
                         [("Hedefin altında", D("21.00")), ("Hedefte", D(0)), ("Hedefte", D(0)), ("Hedefte", D(0)), ("Hedefin üzerinde", D(0))])
        self.assertEqual((d(899, tolerans=D(0))["durum"], d(899, tolerans=D(0))["kayip"], main.analiz([p("P", 1000, 500)])["partiler"][0]["durum"], main.analiz([p("P", 1000, 500)], {}, D(60))["partiler"][0]["sapma"]),
                         ("Hedefin altında", D("1.00"), "Hedef yok", D("-10.00")))

    def test_uyarilar(self):
        t = lambda *a, **kw: sorted(u["tur"] for u in main.analiz([p("P", *a)], {"u": D(90)}, **kw)["uyarilar"])   # noqa: E731
        self.assertEqual((t(1000, 900, 50), t(1000, 900, 49), t(1000, 900, 49, fark_esik=D(6)), t(1000, 1010), t(1000, 950, 60), t(1000, 800, 200)),
                         ([], ["Kayıt dışı fark yüksek"], [], ["Verim %100'ün üzerinde"], ["Çıkış toplamı girişten fazla"], ["Hedefin altında"]))

    def test_ozetler(self):
        r = main.analiz([p("A1", 1000, 925, 50, urun="Bisküvi", hat="Hat 1"), p("A2", 1000, 880, 70, urun="BİSKÜVİ", hat="Hat 1"), p("A3", 1200, 1110, 60, 20, urun="Bisküvi", hat="Hat 2"),
                         p("K1", 800, 690, 30, urun="Kraker", hat="Hat 2")], {"biskuvi": D(92), "kraker": D(90)})
        self.assertEqual([(o["ad"], o["parti"], o["giris"], o["cikis"], o["fire"], o["verim"], o["fire_y"], o["en_dusuk"], o["en_yuksek"], o["hedef"], o["altinda"], o["kayip"]) for o in r["urunler"]],
                         [("Bisküvi", 3, D(3200), D(2915), D(180), D("91.09"), D("5.63"), D("88.00"), D("92.50"), D(92), 1, D("40.00")),
                          ("Kraker", 1, D(800), D(690), D(30), D("86.25"), D("3.75"), D("86.25"), D("86.25"), D(90), 1, D("30.00"))])
        self.assertEqual([(o["ad"], o["parti"], o["giris"], o["verim"], o["altinda"], o["hedef"]) for o in r["hatlar"]], [("Hat 1", 2, D(2000), D("90.25"), 1, None), ("Hat 2", 2, D(2000), D("90.00"), 1, None)])
        g = r["genel"]
        self.assertEqual((g["parti"], g["giris"], g["cikis"], g["fire"], g["rework"], g["verim"], g["fire_y"], g["altinda"], g["kayip"], r["say"]),
                         (4, D(4000), D(3605), D(210), D(20), D("90.13"), D("5.25"), 2, D("70.00"), {"Hedefte": 2, "Hedefin altında": 2}))
        with self.assertRaises(ValueError):
            main.analiz([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "v.xlsx"
            r = main.calistir(ORNEK / "partiler.csv", cikti, ORNEK / "hedefler.csv")
            self.assertEqual([(x["parti"], x["verim"], x["sapma"], x["durum"], x["fark"]) for x in r["partiler"]],
                             [("B-001", D("92.50"), D("0.50"), "Hedefte", D(25)), ("B-002", D("88.00"), D("-4.00"), "Hedefin altında", D(50)), ("B-003", D("92.50"), D("0.50"), "Hedefte", D(10)),
                              ("K-001", D("91.00"), D("1.00"), "Hedefte", D(32)), ("K-002", D("86.25"), D("-3.75"), "Hedefin altında", D(80)), ("K-003", D("102.00"), D("12.00"), "Hedefin üzerinde", D(-15))])
            self.assertEqual([(o["ad"], o["verim"], o["fire_y"]) for o in r["urunler"] + r["hatlar"] + [r["genel"]]],
                             [("Örnek Bisküvi", D("91.09"), D("5.63")), ("Örnek Kraker", D("91.81"), D("3.57")), ("Hat 1", D("92.60"), D("5.00")), ("Hat 2", D("90.29"), D("4.64")), ("Genel", D("91.38"), D("4.81"))])
            self.assertEqual([(u["tur"], u["kim"]) for u in r["uyarilar"]], [("Verim %100'ün üzerinde", "K-003"), ("Hedefin altında", "B-002"), ("Hedefin altında", "K-002"), ("Kayıt dışı fark yüksek", "K-002")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Partiler", "Ürün Özeti", "Hat Özeti", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, h = Path(tmp) / "p.csv", Path(tmp) / "h.csv"
            y.write_text("Parti;Ürün;Giriş;Çıkış;Fire\nP1;A;1.250,5;1100,25;\nP2;A;0;10;\nP3;;100;x;\nP4;B;100;90;-1\nP5;;100;95;2,5\n", encoding="utf-8")
            h.write_text("Ürün;Hedef\nA;%88\nB;120\n", encoding="utf-8")
            l, hedef, uy = main.oku(y, h)
            self.assertEqual(([(x["parti"], x["urun"], x["hat"], x["giris"], x["cikis"], x["fire"]) for x in l], hedef, sorted(u["tur"] for u in uy)),
                             ([("P1", "A", "Belirtilmemiş", D("1250.5"), D("1100.25"), D(0)), ("P5", "Belirtilmemiş", "Belirtilmemiş", D(100), D(95), D("2.5"))], {"a": D(88)}, ["Hedef okunamadı"] + ["Parti okunamadı"] * 3))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--partiler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(arg + ["--hedefler", str(h), "--hedef", "90", "--tolerans", "1", "--fark", "3"]), 0)
            for ek in (["--hedef", "0"], ["--hedef", "x"], ["--tolerans", "-1"], ["--fark", "x"]):
                self.assertEqual(main.main(arg + ek), 2)
            self.assertEqual(main.main(arg + ["--hedefler", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Parti;Ürün\nP1;A\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
