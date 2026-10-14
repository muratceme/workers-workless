import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 14)


def v(ad, tur, gun, maliyet=None, adet=None, kul=None, oto=False, sorumlu="S", pb="TRY"):
    return {"ad": ad, "tur": tur, "saglayici": "", "bitis": BUGUN + timedelta(days=gun), "oto": oto, "sorumlu": sorumlu, "maliyet": None if maliyet is None else D(maliyet),
            "adet": None if adet is None else D(adet), "kullanilan": None if kul is None else D(kul), "pb": pb}


class AnalizTesti(unittest.TestCase):
    def test_durum(self):
        l = [v("a", "SSL sertifikası", -1), v("b", "SSL sertifikası", 7), v("c", "Alan adı", 30), v("d", "Alan adı", 31), v("e", "Yazılım lisansı", 31, 100, 10, 12),
             v("f", "Yazılım lisansı", 91, 50, 10, 5, sorumlu=""), v("g", "Yazılım lisansı", 400, 70, pb="USD")]
        r = main.analiz(l, BUGUN)
        d = {x["ad"]: x["durum"] for x in r["varliklar"]}
        self.assertEqual(d, {"a": "Süresi doldu", "b": "Kritik", "c": "Yaklaşıyor", "d": "İzlemede", "e": "Yaklaşıyor", "f": "Geçerli", "g": "Geçerli"})
        self.assertEqual([x["ad"] for x in r["varliklar"]], ["a", "b", "c", "e", "d", "f", "g"])
        self.assertEqual(r["takvim"], [{"ay": "2026-11", "pb": "TRY", "tutar": D(100)}, {"ay": "2027-01", "pb": "TRY", "tutar": D(50)}])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"][0]) for u in r["uyarilar"]),
                         [("Bilgi", "Atıl lisans", "f"), ("Orta", "Bitiş yaklaşıyor", "c"), ("Orta", "Bitiş yaklaşıyor", "e"), ("Yüksek", "Bitişe çok az kaldı", "b"),
                          ("Yüksek", "Lisans adedi aşıldı", "e"), ("Yüksek", "Süresi doldu", "a")])
        self.assertEqual(main.analiz([v("e", "Yazılım lisansı", 31, adet=10, kul=12)], BUGUN)["varliklar"][0]["kullanim"], D("120.00"))
        with self.assertRaises(ValueError):
            main.analiz([], BUGUN)

    def test_tur(self):
        self.assertEqual([main._tur(x) for x in ("TLS", "Domain", "Bakım sözleşmesi", "SaaS", "", "Donanım")],
                         ["SSL sertifikası", "Alan adı", "Destek / bakım", "Abonelik (SaaS)", "Yazılım lisansı", "Donanım"])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "varliklar.csv", cikti, BUGUN)
            self.assertEqual(r["ozet"], {"Süresi doldu": 1, "Kritik": 2, "Yaklaşıyor": 5, "İzlemede": 0, "Geçerli": 4})
            self.assertEqual(r["varliklar"][0]["ad"], "b2b.ornekfirma.com.tr")
            self.assertEqual(load_workbook(cikti).sheetnames, ["Varlıklar", "Yenileme Takvimi", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "v.csv"
            y.write_text("Ürün;Bitiş;Maliyet;Adet;Doviz\nA;01.11.2026;1.000;5;usd\nB;x;;;\nC;01.11.2026;-1;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["ad"], x["maliyet"], x["pb"], x["tur"]) for x in l], len(uy)), ([("A", D(1000), "USD", "Yazılım lisansı")], 2))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--varliklar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "14.10.2026", "--kritik", "3", "--yakin", "14", "--izle", "60", "--atil", "50"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + x) for x in (["--tarih", "x"], ["--kritik", "40"], ["--atil", "x"])], [2] * 3)
            self.assertEqual(main.main(["--varliklar", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Ürün;Bitiş\nA;x\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
