"""
Grup Rezervasyon Teklifi — Workers / Workless kod bloğu
Rezervasyon ve Gelir Yönetimi › Rezervasyon Sorumlusu
Grup talebini (gün × kalem × adet) otelin fiyat listesiyle fiyatlandırıp teklif dosyası hazırlar.
  Satır tutarı      = adet × o tarihte geçerli KDV hariç fiyat (fiyat listesinde Başlangıç / Bitiş varsa sezon fiyatı)
  Ücretsiz oda      = her gece ⌊oda adedi ÷ (N + 1)⌋ oda, o gecenin en düşük oda fiyatından düşülür (--ucretsiz N; 0 = kapalı)
  İndirim           = (tutar − ücretsiz oda) × --indirim %
  Konaklama vergisi = konaklama vergisine tabi satırların indirimli tutarı × oran (KDV matrahına girmez)
                      varsayılan oran: 01.05.2026–31.12.2026 gecelerine %1, diğerlerine %2 (--kv ile tek oran verilebilir)
  KDV               = her satırın indirimli tutarı × satırın KDV oranı (fiyat listesinden)
  Genel toplam      = net + konaklama vergisi + KDV; kapora = genel toplam × --kapora %
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
FIYAT = {"kalem": ("kalem", "hizmet", "urun", "aciklama"), "tur": ("tur", "kategori", "grup"), "birim": ("birim",), "fiyat": ("fiyat", "birim fiyat", "kdv haric fiyat", "net fiyat"),
         "kdv": ("kdv", "kdv orani", "kdv yuzde"), "kv": ("konaklama vergisi", "kv"), "bas": ("baslangic", "baslangic tarihi", "gecerlilik baslangic"), "bit": ("bitis", "bitis tarihi", "gecerlilik bitis")}
TALEP = {"tarih": ("tarih", "gun"), "kalem": ("kalem", "hizmet", "urun"), "adet": ("adet", "miktar", "oda sayisi", "kisi")}
SIFIR = Decimal(0)


def kv_orani(gun: date) -> Decimal:
    """11263 sayılı Cumhurbaşkanı Kararı: 01.05.2026–31.12.2026 arası %1; diğer dönemler kanuni oran %2 (güncel mevzuattan doğrulayın)."""
    return Decimal(1) if date(2026, 5, 1) <= gun <= date(2026, 12, 31) else Decimal(2)


def fiyat_oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, FIYAT, ("kalem", "fiyat")):
        ad = metin(r["kalem"])
        if not ad:
            continue
        f, kdv = para(r["fiyat"]), para(r["kdv"]) if metin(r["kdv"]) else SIFIR
        bas, bit = (tarih(r[a]) if metin(r[a]) else None for a in ("bas", "bit"))
        if f is None or f < 0 or kdv is None or (metin(r["bas"]) and bas is None) or (metin(r["bit"]) and bit is None):
            uy.append(uyari("Yüksek", "Fiyat satırı okunamadı", ad, f"Satır {r['_satir']}: fiyat, KDV ya da geçerlilik tarihi okunamadı."))
            continue
        tur = metin(r["tur"]) or "Diğer"
        kv = katla(r["kv"])
        l.append({"kalem": ad, "tur": tur, "birim": metin(r["birim"]), "fiyat": f, "kdv": kdv, "bas": bas, "bit": bit,
                  "kv": (kv in ("evet", "e", "var", "x", "1")) if kv else katla(tur) == "oda", "oda": katla(tur) == "oda"})
    return l, uy


def talep_oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, TALEP, ("tarih", "kalem", "adet")):
        if not metin(r["kalem"]) and not metin(r["tarih"]):
            continue
        g, a = tarih(r["tarih"]), para(r["adet"])
        if g is None or a is None or a <= 0 or not metin(r["kalem"]):
            uy.append(uyari("Yüksek", "Talep satırı okunamadı", metin(r["kalem"]) or f"Satır {r['_satir']}", f"Satır {r['_satir']}: tarih, kalem ya da adet okunamadı."))
            continue
        l.append({"tarih": g, "kalem": metin(r["kalem"]), "adet": a})
    return l, uy


def hesapla(talep: list[dict], fiyatlar: list[dict], ucretsiz: int = 0, indirim: Decimal = SIFIR, kv: Decimal | None = None, kisi: int | None = None,
            kapora: Decimal = Decimal(30), teklif: date | None = None, opsiyon: int = 7, bakiye_gun: int = 7) -> dict:
    if not talep:
        raise ValueError("Talep satırı yok.")
    uy, satirlar = [], []
    for t in sorted(talep, key=lambda t: t["tarih"]):
        aday = [f for f in fiyatlar if katla(f["kalem"]) == katla(t["kalem"]) and (f["bas"] is None or f["bas"] <= t["tarih"]) and (f["bit"] is None or t["tarih"] <= f["bit"])]
        if not aday:
            uy.append(uyari("Yüksek", "Fiyat bulunamadı", t["kalem"], f"{t['tarih']:%d.%m.%Y} için fiyat listesinde geçerli fiyat yok; satır teklife girmedi."))
            continue
        if len(aday) > 1:
            uy.append(uyari("Orta", "Birden fazla fiyat", t["kalem"], f"{t['tarih']:%d.%m.%Y} için {len(aday)} geçerli fiyat var; ilki ({aday[0]['fiyat']:,.2f}) kullanıldı."))
        f = aday[0]
        satirlar.append({"tarih": t["tarih"], "kalem": f["kalem"], "tur": f["tur"], "birim": f["birim"], "adet": t["adet"], "fiyat": f["fiyat"], "tutar": yuvarla(t["adet"] * f["fiyat"]),
                         "kdv_oran": f["kdv"], "kv": f["kv"], "oda": f["oda"], "kv_oran": kv if kv is not None else kv_orani(t["tarih"])})
    if not satirlar:
        raise ValueError("Fiyatlandırılabilen talep satırı yok.")
    if ucretsiz:
        for g in sorted({s["tarih"] for s in satirlar}):
            odalar = [s for s in satirlar if s["tarih"] == g and s["oda"]]
            n = int(sum(s["adet"] for s in odalar) // (ucretsiz + 1))
            if n:
                en_ucuz = min(odalar, key=lambda s: s["fiyat"])
                satirlar.append({"tarih": g, "kalem": f"Ücretsiz oda ({ucretsiz}+1) · {en_ucuz['kalem']}", "tur": "Oda", "birim": en_ucuz["birim"], "adet": Decimal(-n), "fiyat": en_ucuz["fiyat"],
                                 "tutar": -yuvarla(n * en_ucuz["fiyat"]), "kdv_oran": en_ucuz["kdv_oran"], "kv": en_ucuz["kv"], "oda": True, "kv_oran": en_ucuz["kv_oran"], "ucretsiz": True})
        satirlar.sort(key=lambda s: (s["tarih"], s.get("ucretsiz", False)))
    for s in satirlar:
        s["indirim"] = yuvarla(s["tutar"] * indirim / 100)
        s["net"] = s["tutar"] - s["indirim"]
        s["kv_tutar"] = yuvarla(s["net"] * s["kv_oran"] / 100) if s["kv"] else SIFIR
        s["kdv"] = yuvarla(s["net"] * s["kdv_oran"] / 100)
        s["toplam"] = s["net"] + s["kv_tutar"] + s["kdv"]
    t = {a: sum((s[a] for s in satirlar), SIFIR) for a in ("tutar", "indirim", "net", "kv_tutar", "kdv", "toplam")}
    t["brut"] = sum((s["tutar"] for s in satirlar if not s.get("ucretsiz")), SIFIR)
    t["ucretsiz"] = -sum((s["tutar"] for s in satirlar if s.get("ucretsiz")), SIFIR)
    t["ucretsiz_adet"] = -int(sum((s["adet"] for s in satirlar if s.get("ucretsiz")), SIFIR))
    oda_gece = int(sum((s["adet"] for s in satirlar if s["oda"] and not s.get("ucretsiz")), SIFIR))
    oda_net = sum((s["net"] for s in satirlar if s["oda"]), SIFIR)
    t.update({"oda_gece": oda_gece, "ort_oda": yuvarla(oda_net / oda_gece) if oda_gece else None, "kisi": kisi, "kisi_basi": yuvarla(t["toplam"] / kisi) if kisi else None,
              "kapora": yuvarla(t["toplam"] * kapora / 100), "kapora_oran": kapora, "ilk": satirlar[0]["tarih"], "son": max(s["tarih"] for s in satirlar)})
    t["opsiyon"] = teklif + timedelta(days=opsiyon) if teklif else None
    t["bakiye_tarih"] = t["ilk"] - timedelta(days=bakiye_gun)
    if teklif and t["bakiye_tarih"] < teklif:
        uy.append(uyari("Orta", "Bakiye tarihi geçmiş", f"{t['bakiye_tarih']:%d.%m.%Y}", f"Grup girişe {(t['ilk'] - teklif).days} gün kala teklif ediliyor; bakiye ödeme tarihi teklif tarihinden önce kalıyor, ödeme planını elle belirleyin."))
    if teklif and t["opsiyon"] and t["opsiyon"] >= t["ilk"]:
        uy.append(uyari("Orta", "Opsiyon girişten sonra", f"{t['opsiyon']:%d.%m.%Y}", "Opsiyon tarihi grubun giriş tarihine denk geliyor ya da geçiyor; daha kısa opsiyon verin."))
    if kisi is None:
        uy.append(uyari("Bilgi", "Kişi sayısı verilmedi", "--kisi", "Kişi başı fiyat hesaplanmadı."))
    kalem = {}
    for s in satirlar:
        k = kalem.setdefault(s["kalem"], {"kalem": s["kalem"], "tur": s["tur"], "birim": s["birim"], "adet": SIFIR, "net": SIFIR, "toplam": SIFIR, "gun": set()})
        k["adet"] += s["adet"]
        k["net"] += s["net"]
        k["toplam"] += s["toplam"]
        k["gun"].add(s["tarih"])
    return {"satirlar": satirlar, "kalemler": list(kalem.values()), "toplam": t, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, grup: str, teklif: date | None) -> None:
    t = r["toplam"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Teklif"
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 38, 20
    ws.append(["GRUP REZERVASYON TEKLİFİ"])
    ws["A1"].font = Font(bold=True, size=14)
    for a, d in (("Grup", grup), ("Teklif tarihi", teklif), ("Giriş", t["ilk"]), ("Son hizmet günü", t["son"]), ("Kişi", t["kisi"]), ("Toplam oda-gece (ücretli + ücretsiz)", t["oda_gece"])):
        _satir(ws, [a, d], tarih_s=(2,) if isinstance(d, date) else ())
    ws.append([])
    _baslik(ws, ["Tutar", "TL"], dondur=None)
    for a, d in (("Hizmet tutarı (KDV hariç)", t["brut"]), (f"Ücretsiz oda ({t['ucretsiz_adet']} oda-gece)", -t["ucretsiz"]), ("Grup indirimi", -t["indirim"]),
                 ("Net tutar (KDV hariç)", t["net"]), ("Konaklama vergisi", t["kv_tutar"]), ("KDV", t["kdv"]), ("GENEL TOPLAM", t["toplam"]),
                 ("Ortalama net oda fiyatı (oda-gece)", t["ort_oda"]), ("Kişi başı (genel toplam ÷ kişi)", t["kisi_basi"])):
        n = _satir(ws, [a, d], para_s=(2,))
        if a == "GENEL TOPLAM":
            ws.cell(n, 1).font = ws.cell(n, 2).font = Font(bold=True)
    ws.append([])
    _baslik(ws, ["Koşul", ""], dondur=None)
    kosul = [("Opsiyon (teklif geçerlilik) tarihi", t["opsiyon"]), (f"Kapora (%{t['kapora_oran']})", t["kapora"]), ("Bakiye son ödeme tarihi", t["bakiye_tarih"])]
    for a, d in kosul:
        _satir(ws, [a, d], para_s=(2,) if isinstance(d, Decimal) else (), tarih_s=(2,) if isinstance(d, date) else ())
    ws.append([])
    for not_ in ("Fiyatlar TL'dir; KDV ve konaklama vergisi ayrıca gösterilmiştir.", "Teklif, opsiyon tarihine kadar müsaitliğe bağlı olarak geçerlidir; kesinleşme kapora ile olur.",
                 "Oda sayısı ve kişi sayısındaki değişiklikler, iptal ve düşüş (attrition) koşulları sözleşmede ayrıca belirlenir."):
        ws.append([not_])
    ws = wb.create_sheet("Detay")
    _baslik(ws, ["Tarih", "Kalem", "Tür", "Birim", "Adet", "Birim Fiyat", "Tutar", "İndirim", "Net", "KV %", "Konaklama V.", "KDV %", "KDV", "Toplam"],
            (11, 42, 8, 10, 7, 12, 13, 11, 13, 6, 11, 6, 12, 13))
    for s in r["satirlar"]:
        _satir(ws, [s["tarih"], s["kalem"], s["tur"], s["birim"], s["adet"], s["fiyat"], s["tutar"], s["indirim"], s["net"], s["kv_oran"] if s["kv"] else None, s["kv_tutar"],
                    s["kdv_oran"], s["kdv"], s["toplam"]], para_s=(6, 7, 8, 9, 11, 13, 14), tarih_s=(1,), dolgu={2: YESIL if s.get("ucretsiz") else None})
    ws = wb.create_sheet("Kalem Özeti")
    _baslik(ws, ["Kalem", "Tür", "Birim", "Gün", "Toplam Adet", "Net", "Toplam (vergiler dahil)"], (42, 8, 10, 6, 11, 14, 16))
    for k in r["kalemler"]:
        _satir(ws, [k["kalem"], k["tur"], k["birim"], len(k["gun"]), k["adet"], k["net"], k["toplam"]], para_s=(6, 7))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(talep_yol: Path, fiyat_yol: Path, cikti: Path, grup: str = "", **k) -> dict:
    fiyatlar, uy1 = fiyat_oku(fiyat_yol)
    talep, uy2 = talep_oku(talep_yol)
    r = hesapla(talep, fiyatlar, **k)
    r["uyarilar"] = sirala(uy1 + uy2 + r["uyarilar"])
    rapor_yaz(cikti, r, grup, k.get("teklif"))
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Grup talebini fiyat listesiyle fiyatlandırıp teklif dosyası hazırlar.")
    p.add_argument("--talep", type=Path, help="Tarih, Kalem, Adet")
    p.add_argument("--fiyatlar", type=Path, help="Kalem, Tür (Oda / Yemek / Salon / Diğer), Birim, Fiyat (KDV hariç), KDV %%, Konaklama Vergisi (Evet / Hayır), Başlangıç, Bitiş")
    p.add_argument("--grup", default="", help="Grup / firma adı")
    p.add_argument("--kisi", type=int, help="Grup kişi sayısı (kişi başı fiyat için)")
    p.add_argument("--ucretsiz", type=int, default=0, help="Her N ücretli odaya 1 ücretsiz oda (N+1); 0 = uygulanmaz")
    p.add_argument("--indirim", default="0", help="Grup indirimi yüzdesi (tüm satırlara)")
    p.add_argument("--kv", help="Konaklama vergisi oranı; verilmezse tarihe göre (01.05–31.12.2026 için 1, diğer 2)")
    p.add_argument("--kapora", default="30", help="Kapora yüzdesi (varsayılan 30)")
    p.add_argument("--teklif-tarihi", help="GG.AA.YYYY (varsayılan bugün; örnek veride 14.10.2026)")
    p.add_argument("--opsiyon", type=int, default=7, help="Opsiyon süresi, gün (varsayılan 7)")
    p.add_argument("--bakiye-gun", type=int, default=7, help="Bakiye girişten kaç gün önce ödenir (varsayılan 7)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "grup_teklifi.xlsx")
    a = p.parse_args(argv)
    talep, fiyat = a.talep or ORNEK / "talep.csv", a.fiyatlar or ORNEK / "fiyatlar.csv"
    ornek = a.talep is None and a.fiyatlar is None
    teklif = tarih(a.teklif_tarihi) if a.teklif_tarihi else (date(2026, 10, 14) if ornek else date.today())
    ind, kap, kv = para(a.indirim), para(a.kapora), para(a.kv) if a.kv else None
    if (teklif is None or ind is None or not 0 <= ind < 100 or kap is None or not 0 <= kap <= 100 or (a.kv and (kv is None or not 0 <= kv <= 100))
            or a.ucretsiz < 0 or a.opsiyon < 0 or a.bakiye_gun < 0 or (a.kisi is not None and a.kisi < 1)):
        print("[X] --indirim 0-99, --kapora / --kv 0-100, --ucretsiz / --opsiyon / --bakiye-gun 0 ya da üstü, --kisi en az 1, --teklif-tarihi GG.AA.YYYY olmalı")
        return 2
    for y in (talep, fiyat):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(talep, fiyat, a.cikti, a.grup or ("Örnek Kongre Grubu" if ornek else ""), ucretsiz=a.ucretsiz or (20 if ornek else 0), indirim=ind, kv=kv,
                     kisi=a.kisi or (48 if ornek else None), kapora=kap, teklif=teklif, opsiyon=a.opsiyon, bakiye_gun=a.bakiye_gun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {t['ilk']:%d.%m.%Y}–{t['son']:%d.%m.%Y} · {t['oda_gece']} oda-gece · net {t['net']:,.2f} · KV {t['kv_tutar']:,.2f} · KDV {t['kdv']:,.2f} · genel toplam {t['toplam']:,.2f}"
          + (f" · kişi başı {t['kisi_basi']:,.2f}" if t["kisi_basi"] is not None else ""))
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Teklif: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
