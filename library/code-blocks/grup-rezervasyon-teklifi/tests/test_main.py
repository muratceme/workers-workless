import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def f(kalem, fiyat, tur="Oda", kdv=10, kv=None, bas=None, bit=None):
    return {"kalem": kalem, "tur": tur, "birim": "", "fiyat": D(fiyat), "kdv": D(kdv), "bas": bas, "bit": bit, "kv": (tur == "Oda") if kv is None else kv, "oda": tur == "Oda"}


def t(g, kalem, adet):
    return {"tarih": date(2026, 11, g), "kalem": kalem, "adet": D(adet)}


class HesapTesti(unittest.TestCase):
    def test_vergiler_ve_ucretsiz(self):
        fy = [f("STD", 1000), f("DLX", 1500), f("Salon", 10000, "Salon", 20), f("Yemek", 500, "Yemek")]
        r = main.hesapla([t(1, "STD", 20), t(1, "DLX", 1), t(1, "Salon", 1), t(1, "Yemek", 30)], fy, ucretsiz=20, indirim=D(10), kisi=30, teklif=date(2026, 10, 1))
        x = r["toplam"]
        # brüt 20000+1500+10000+15000 = 46500; ücretsiz 1 × 1000; indirim %10 → net 40950
        self.assertEqual((x["brut"], x["ucretsiz"], x["ucretsiz_adet"], x["indirim"], x["net"]), (D(46500), D(1000), 1, D(4550), D(40950)))
        # oda net (20000 + 1500 − 1000) × 0,9 = 18450 → KV %1 = 184,50; KDV 18450×%10 + 9000×%20 + 13500×%10 = 1845 + 1800 + 1350
        self.assertEqual((x["kv_tutar"], x["kdv"], x["toplam"], x["oda_gece"], x["ort_oda"], x["kisi_basi"], x["kapora"]),
                         (D("184.50"), D(4995), D("46129.50"), 21, D("878.57"), D("1537.65"), D("13838.85")))
        self.assertEqual((x["opsiyon"], x["bakiye_tarih"]), (date(2026, 10, 8), date(2026, 10, 25)))
        self.assertEqual(r["uyarilar"], [])

    def test_sezon_ve_uyarilar(self):
        fy = [f("STD", 1000, bit=date(2026, 11, 1)), f("STD", 800, bas=date(2026, 11, 2)), f("X", 5, "Diğer"), f("X", 6, "Diğer")]
        r = main.hesapla([t(1, "std", 2), t(2, "STD", 2), t(2, "X", 1), t(2, "Y", 1)], fy, kv=D(2), teklif=date(2026, 10, 30), opsiyon=5)
        self.assertEqual([s["fiyat"] for s in r["satirlar"]], [D(1000), D(800), D(5)])
        self.assertEqual(r["toplam"]["kv_tutar"], D(72))
        self.assertEqual(sorted(u["tur"] for u in r["uyarilar"]), ["Bakiye tarihi geçmiş", "Birden fazla fiyat", "Fiyat bulunamadı", "Kişi sayısı verilmedi", "Opsiyon girişten sonra"])
        self.assertEqual([main.kv_orani(g) for g in (date(2026, 4, 30), date(2026, 5, 1), date(2026, 12, 31), date(2027, 1, 1))], [D(2), D(1), D(1), D(2)])
        with self.assertRaises(ValueError):
            main.hesapla([t(1, "Y", 1)], fy)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "talep.csv", ORNEK / "fiyatlar.csv", cikti, "Grup", ucretsiz=20, kisi=48, teklif=date(2026, 10, 14))
            x = r["toplam"]
            self.assertEqual((x["brut"], x["ucretsiz"], x["net"], x["kv_tutar"], x["kdv"], x["toplam"], x["oda_gece"], x["ort_oda"], x["kisi_basi"], x["kapora"]),
                             (D(488140), D(9300), D(478840), D(2233), D(51934), D(533007), 71, D("3145.07"), D("11104.31"), D("159902.10")))
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]], [("Yüksek", "Fiyat bulunamadı", "Transfer (otobüs)")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Teklif", "Detay", "Kalem Özeti", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            fy, ty = Path(tmp) / "f.csv", Path(tmp) / "t.csv"
            fy.write_text("Hizmet;Kategori;Birim Fiyat;KDV Oranı;Başlangıç Tarihi\nA;Oda;1.000;10;\nB;Yemek;x;10;\nC;Salon;5;;01.13.2026\nD;Yemek;10;;\n", encoding="utf-8")
            ty.write_text("Gün;Hizmet;Miktar\n01.11.2026;A;2\n01.11.2026;D;0\nx;A;1\n", encoding="utf-8")
            l, uy = main.fiyat_oku(fy)
            self.assertEqual(([(x["kalem"], x["fiyat"], x["kv"]) for x in l], len(uy)), ([("A", D(1000), True), ("D", D(10), False)], 2))
            l, uy = main.talep_oku(ty)
            self.assertEqual((len(l), len(uy)), (1, 2))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--talep", str(ty), "--fiyatlar", str(fy), "--cikti", c]
            self.assertEqual(main.main(arg + ["--kisi", "2", "--indirim", "5", "--kv", "2", "--teklif-tarihi", "01.10.2026"]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--indirim", "100"], ["--kapora", "x"], ["--kv", "101"], ["--ucretsiz", "-1"], ["--kisi", "0"], ["--teklif-tarihi", "x"])], [2] * 6)
            self.assertEqual(main.main(["--talep", str(Path(tmp) / "yok.csv"), "--cikti", c]), 1)
            ty.write_text("Gün;Hizmet;Miktar\n01.11.2026;Z;2\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
