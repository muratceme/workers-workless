"""
Sevkiyat Araç Planlama — Gerisi Bende kod bloğu
Lojistik ve Depo › Lojistik Uzmanı

Siparişleri teslim tarihi ve bölgeye göre gruplar; palet, ağırlık ve hacim kapasitelerine göre araçlara yerleştirir,
araç tipini seçer ve doluluk planı çıkarır.
  1. Grup: aynı teslim tarihi + aynı bölge.
  2. Tek araca sığmayan sipariş, en büyük aracın kapasitesi kadar parçalara bölünür (uyarı verilir).
  3. Yerleştirme: siparişler büyükten küçüğe sıralanır. Her adımda her araç tipi için bir araç doldurulur (sığan ilk
     sipariş eklenir); taşınan yük başına maliyeti en düşük olan tip seçilir. Siparişler bitene kadar tekrarlanır.
  4. Küçültme: her sefer, yükünü taşıyabilen en ucuz araç tipine indirilir.
  5. Aynı grup yalnız tek araç tipiyle de denenir (her tip için ayrı); toplam maliyeti en düşük plan seçilir.
  Doluluk = palet, ağırlık ve hacim doluluklarının en yükseği (aracı sınırlayan boyut).
Araç kapasiteleri ve sefer maliyetleri sizin araç dosyanızdan okunur; kodda sabit kapasite yoktur.
Bu bir sezgisel yöntemdir; en iyi çözümü garanti etmez. Rota, yükleme sırası ve dingil ağırlığı hesaplanmaz.
İnternete bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: 3 gün, 4 bölge, 34 sipariş
    python main.py --siparisler siparisler.xlsx --araclar araclar.xlsx --min-doluluk 60
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SIPARIS_SUTUNLARI = {"no": ("siparis no", "siparis", "irsaliye no"), "musteri": ("musteri", "alici", "cari"), "bolge": ("teslim bolgesi", "bolge", "il", "varis"),
                     "tarih": ("teslim tarihi", "sevk tarihi", "tarih"), "palet": ("palet", "palet sayisi", "palet adedi"),
                     "kg": ("agirlik", "agirlik kg", "brut agirlik", "kg"), "m3": ("hacim", "hacim m3", "m3")}
ARAC_SUTUNLARI = {"tip": ("arac tipi", "tip", "arac"), "palet": ("palet kapasitesi", "palet"), "kg": ("agirlik kapasitesi", "agirlik kapasitesi kg", "tasima kapasitesi kg", "kg"),
                  "m3": ("hacim kapasitesi", "hacim kapasitesi m3", "hacim", "m3"), "maliyet": ("sefer maliyeti", "sefer maliyeti tl", "maliyet"),
                  "adet": ("mevcut adet", "adet", "arac sayisi"), "bolge": ("bolge", "bolge maliyeti icin")}
BOYUTLAR = (("palet", "Palet"), ("kg", "Ağırlık"), ("m3", "Hacim"))


@dataclass
class Siparis:
    no: str
    musteri: str
    bolge: str
    tarih: date
    palet: Decimal
    kg: Decimal
    m3: Decimal
    parca: str = ""


@dataclass
class Arac:
    tip: str
    palet: Decimal
    kg: Decimal
    m3: Decimal
    maliyet: Decimal
    adet: int | None = None          # None = sınırsız


@dataclass
class Sefer:
    tarih: date
    bolge: str
    arac: Arac
    siparisler: list = field(default_factory=list)

    def yuk(self, boyut: str) -> Decimal:
        return sum((getattr(s, boyut) for s in self.siparisler), Decimal(0))

    def doluluk(self) -> tuple[Decimal, str]:
        oranlar = [(self.yuk(b) / getattr(self.arac, b) * 100, ad) for b, ad in BOYUTLAR if getattr(self.arac, b)]
        return max(oranlar) if oranlar else (Decimal(0), "")


def oku(s_yol: Path, a_yol: Path) -> tuple[list[Siparis], list[Arac], list[dict]]:
    uy, sip, arac = [], [], []
    for r in kayitlar(a_yol, ARAC_SUTUNLARI, ("tip", "maliyet")):
        if not metin(r["tip"]):
            continue
        p, k, m, c = (para(r[x]) or Decimal(0) for x in ("palet", "kg", "m3", "maliyet"))
        if c <= 0 or not (p or k or m):
            uy.append(uyari("Yüksek", "Araç tanımı eksik", r["tip"], f"Satır {r['_satir']}: sefer maliyeti ve en az bir kapasite girilmeli; araç tipi kullanılmadı."))
            continue
        adet = para(r["adet"])
        arac.append(Arac(metin(r["tip"]), p, k, m, c, int(adet) if adet is not None else None))
    if not arac:
        raise ValueError(f"{a_yol.name}: kullanılabilir araç tipi yok")
    for r in kayitlar(s_yol, SIPARIS_SUTUNLARI, ("no", "tarih")):
        if not metin(r["no"]):
            continue
        t = tarih(r["tarih"])
        p, k, m = (para(r[x]) or Decimal(0) for x in ("palet", "kg", "m3"))
        if not t:
            uy.append(uyari("Yüksek", "Teslim tarihi okunamadı", r["no"], f"Satır {r['_satir']}: '{metin(r['tarih'])}'. Sipariş plana alınmadı."))
            continue
        if p < 0 or k < 0 or m < 0 or not (p or k or m):
            uy.append(uyari("Yüksek", "Yük bilgisi yok", r["no"], f"Satır {r['_satir']}: palet, ağırlık ve hacmin en az biri sıfırdan büyük olmalı. Sipariş plana alınmadı."))
            continue
        eksik = [ad for b, ad in BOYUTLAR if not {"palet": p, "kg": k, "m3": m}[b]]
        if eksik:
            uy.append(uyari("Bilgi", "Eksik boyut", r["no"], f"{', '.join(eksik)} girilmemiş; bu boyut kapasite kontrolüne katılmadı."))
        sip.append(Siparis(metin(r["no"]), metin(r["musteri"]), metin(r["bolge"]) or "Belirtilmemiş", t, p, k, m))
    return sip, arac, uy


def sigar(sefer_yuk: dict, s: Siparis, a: Arac) -> bool:
    """Kapasitesi 0 girilen boyut o araç için kontrol edilmez."""
    return all(not getattr(a, b) or sefer_yuk[b] + getattr(s, b) <= getattr(a, b) for b, _ in BOYUTLAR)


def buyukluk(s, ref: dict) -> Decimal:
    return max((getattr(s, b) / ref[b] for b, _ in BOYUTLAR if ref[b]), default=Decimal(0))


def bol(s: Siparis, araclar: list[Arac], uy: list[dict]) -> list[Siparis]:
    """Hiçbir araca tek başına sığmayan siparişi, en çok yük alan aracın kapasitesi kadar parçalara böler."""
    if any(sigar({b: Decimal(0) for b, _ in BOYUTLAR}, s, a) for a in araclar):
        return [s]
    en_buyuk = max(araclar, key=lambda a: min((getattr(a, b) / getattr(s, b) for b, _ in BOYUTLAR if getattr(s, b) and getattr(a, b)), default=Decimal(0)))
    oran = min((getattr(en_buyuk, b) / getattr(s, b) for b, _ in BOYUTLAR if getattr(s, b) and getattr(en_buyuk, b)), default=Decimal(1))
    if s.palet:                                                 # palet bölünmez: tam palete yuvarla
        tam = (s.palet * oran).to_integral_value(rounding="ROUND_FLOOR")
        if tam < 1:
            uy.append(uyari("Yüksek", "Sipariş araca sığmıyor", s.no, "Bir paletin ağırlığı veya hacmi en büyük aracın kapasitesini aşıyor. Sipariş plana alınmadı."))
            return []
        oran = tam / s.palet
    parcalar, kalan = [], Siparis(s.no, s.musteri, s.bolge, s.tarih, s.palet, s.kg, s.m3)
    parca = Siparis(s.no, s.musteri, s.bolge, s.tarih, s.palet * oran, yuvarla(s.kg * oran), yuvarla(s.m3 * oran, "0.001"))
    while not sigar({b: Decimal(0) for b, _ in BOYUTLAR}, kalan, en_buyuk):
        parcalar.append(Siparis(s.no, s.musteri, s.bolge, s.tarih, parca.palet, parca.kg, parca.m3))
        kalan = Siparis(s.no, s.musteri, s.bolge, s.tarih, kalan.palet - parca.palet, kalan.kg - parca.kg, kalan.m3 - parca.m3)
    parcalar.append(kalan)
    for i, p in enumerate(parcalar, 1):
        p.parca = f"{i}/{len(parcalar)}"
    uy.append(uyari("Orta", "Sipariş bölündü", s.no, f"Tek araca sığmıyor ({s.palet:f} palet, {s.kg:f} kg, {s.m3:f} m³); {len(parcalar)} parçaya bölündü ({en_buyuk.tip} kapasitesine göre)."))
    return parcalar


def _dene(bekleyen: list[Siparis], tipler: list[Arac], araclar: list[Arac], kalan_adet: dict, ref: dict) -> tuple[list[Sefer], list[Siparis], dict]:
    """Yalnız `tipler` ile yerleştirir, sonra her seferi yükünü taşıyan en ucuz tipe (tüm araçlar içinden) indirir."""
    kalan_adet, bekleyen, seferler = dict(kalan_adet), list(bekleyen), []
    while bekleyen:
        en_iyi = None
        for a in tipler:
            if kalan_adet[a.tip] is not None and kalan_adet[a.tip] <= 0:
                continue
            yuk, alinan = {b: Decimal(0) for b, _ in BOYUTLAR}, []
            for s in bekleyen:
                if sigar(yuk, s, a):
                    alinan.append(s)
                    for b, _ in BOYUTLAR:
                        yuk[b] += getattr(s, b)
            tasinan = sum((buyukluk(s, ref) for s in alinan), Decimal(0))
            if alinan and (en_iyi is None or a.maliyet / tasinan < en_iyi[0]):
                en_iyi = (a.maliyet / tasinan, a, alinan)
        if en_iyi is None:
            break
        _, a, alinan = en_iyi
        if kalan_adet[a.tip] is not None:
            kalan_adet[a.tip] -= 1
        seferler.append(Sefer(alinan[0].tarih, alinan[0].bolge, a, alinan))
        bekleyen = [s for s in bekleyen if not any(s is x for x in alinan)]
    for sf in seferler:
        yuk = {b: sf.yuk(b) for b, _ in BOYUTLAR}
        for a in sorted(araclar, key=lambda a: a.maliyet):
            if a.maliyet >= sf.arac.maliyet:
                break
            if (kalan_adet[a.tip] is None or kalan_adet[a.tip] > 0) and all(not getattr(a, b) or yuk[b] <= getattr(a, b) for b, _ in BOYUTLAR):
                if kalan_adet[sf.arac.tip] is not None:
                    kalan_adet[sf.arac.tip] += 1
                if kalan_adet[a.tip] is not None:
                    kalan_adet[a.tip] -= 1
                sf.arac = a
                break
    return seferler, bekleyen, kalan_adet


def grubu_yerlestir(siparisler: list[Siparis], araclar: list[Arac], kalan_adet: dict, uy: list[dict]) -> list[Sefer]:
    """Karma (tüm tipler) ve tek tipli denemelerden, en az siparişi açıkta bırakan en ucuz planı seçer."""
    ref = {b: max(getattr(a, b) for a in araclar) for b, _ in BOYUTLAR}
    bekleyen = sorted((p for s in siparisler for p in bol(s, araclar, uy)), key=lambda s: (-buyukluk(s, ref), s.no, s.parca))
    if not bekleyen:
        return []
    denemeler = [_dene(bekleyen, tipler, araclar, kalan_adet, ref) for tipler in [araclar] + [[a] for a in araclar]]
    seferler, acikta, yeni_adet = min(denemeler, key=lambda d: (len(d[1]), sum(sf.arac.maliyet for sf in d[0]), len(d[0])))
    kalan_adet.update(yeni_adet)
    for s in acikta:
        uy.append(uyari("Yüksek", "Araç yetersiz", s.no + (f" ({s.parca})" if s.parca else ""),
                        f"{s.tarih:%d.%m.%Y} {s.bolge}: uygun araç kalmadı ({s.palet:f} palet, {s.kg:f} kg). Sipariş atanamadı."))
    return seferler


def planla(siparisler: list[Siparis], araclar: list[Arac], min_doluluk: Decimal = Decimal(60)) -> dict:
    uy, seferler = [], []
    gruplar = defaultdict(list)
    for s in siparisler:
        gruplar[(s.tarih, s.bolge)].append(s)
    gunluk_adet: dict = {}
    for (t, bolge), liste in sorted(gruplar.items()):
        kalan = gunluk_adet.setdefault(t, {a.tip: a.adet for a in araclar})        # araç adedi günlük sayılır
        seferler += grubu_yerlestir(liste, araclar, kalan, uy)
    for i, sf in enumerate(seferler, 1):
        sf.no = f"S{i:03d}"
        d, kisit = sf.doluluk()
        if d < min_doluluk:
            komsu = [x for x in seferler if x is not sf and x.bolge == sf.bolge and abs((x.tarih - sf.tarih).days) == 1]
            uy.append(uyari("Orta", "Düşük doluluk", f"{sf.no} · {sf.tarih:%d.%m.%Y} {sf.bolge}",
                            f"{sf.arac.tip} doluluğu %{yuvarla(d, '0.1')}. " + ("Aynı bölgeye bir gün arayla başka sefer var; teslim tarihi kaydırılabiliyorsa birleştirmeyi değerlendirin."
                                                                               if komsu else "Parsiyel taşıma veya komşu bölgeyle birleştirme değerlendirilebilir.")))
    return {"seferler": seferler, "uyarilar": uy, "araclar": araclar, "siparis": len(siparisler)}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

def rapor_yaz(cikti: Path, r: dict) -> None:
    wb, sf_l = Workbook(), r["seferler"]
    pl = wb.active
    pl.title = "Araç Planı"
    _baslik(pl, ["Sefer", "Tarih", "Bölge", "Araç Tipi", "Sipariş", "Palet", "Palet Kap.", "Ağırlık (kg)", "Ağırlık Kap.", "Hacim (m³)", "Hacim Kap.", "Doluluk %",
                 "Sınırlayan", "Sefer Maliyeti", "Palet Başına Maliyet"], (8, 11, 16, 20, 8, 8, 8, 12, 12, 10, 10, 10, 11, 14, 14))
    for s in sf_l:
        d, kisit = s.doluluk()
        p = s.yuk("palet")
        _satir(pl, [s.no, s.tarih, s.bolge, s.arac.tip, len(s.siparisler), p, s.arac.palet, s.yuk("kg"), s.arac.kg, s.yuk("m3"), s.arac.m3, yuvarla(d, "0.1"), kisit,
                    s.arac.maliyet, yuvarla(s.arac.maliyet / p) if p else None], para_s=(8, 9, 14, 15), tarih_s=(2,),
               dolgu={12: YESIL if d >= 85 else SARI if d >= 60 else KIRMIZI})
    pl.auto_filter.ref = f"A1:O{pl.max_row}"

    yl = wb.create_sheet("Yükleme Listesi")
    _baslik(yl, ["Sefer", "Tarih", "Bölge", "Araç Tipi", "Sipariş No", "Parça", "Müşteri", "Palet", "Ağırlık (kg)", "Hacim (m³)"], (8, 11, 16, 20, 14, 7, 30, 8, 12, 10))
    for s in sf_l:
        for x in s.siparisler:
            _satir(yl, [s.no, s.tarih, s.bolge, s.arac.tip, x.no, x.parca, x.musteri, x.palet, x.kg, x.m3], para_s=(9,), tarih_s=(2,))
    yl.auto_filter.ref = f"A1:J{yl.max_row}"

    oz = wb.create_sheet("Özet")
    _baslik(oz, ["Araç Tipi", "Sefer", "Palet", "Ağırlık (kg)", "Ortalama Doluluk %", "Toplam Maliyet"], (22, 8, 9, 13, 12, 15))
    for a in r["araclar"]:
        l = [s for s in sf_l if s.arac is a]
        if l:
            _satir(oz, [a.tip, len(l), sum(s.yuk("palet") for s in l), sum(s.yuk("kg") for s in l), yuvarla(sum(s.doluluk()[0] for s in l) / len(l), "0.1"),
                        a.maliyet * len(l)], para_s=(4, 6))
    if sf_l:
        _satir(oz, ["Toplam", len(sf_l), sum(s.yuk("palet") for s in sf_l), sum(s.yuk("kg") for s in sf_l),
                    yuvarla(sum(s.doluluk()[0] for s in sf_l) / len(sf_l), "0.1"), sum(s.arac.maliyet for s in sf_l)], para_s=(4, 6))
        for c in oz[oz.max_row]:
            c.font = Font(bold=True)
    oz.append([])
    oz.append(["Doluluk, palet / ağırlık / hacim doluluklarının en yükseğidir. Plan sezgisel bir yöntemle üretilir; rota ve yükleme sırası hesaplanmaz."])
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(s_yol: Path, a_yol: Path, cikti: Path, min_doluluk: Decimal = Decimal(60)) -> dict:
    sip, arac, uy = oku(s_yol, a_yol)
    if not sip:
        raise ValueError(f"{s_yol.name}: planlanabilir sipariş yok")
    r = planla(sip, arac, min_doluluk)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Siparişleri palet, ağırlık ve hacme göre araçlara yerleştirir; araç tipi ve doluluk planı çıkarır.")
    p.add_argument("--siparisler", type=Path, default=ORNEK / "siparisler.csv", help="Sipariş No, Müşteri, Teslim Bölgesi, Teslim Tarihi, Palet, Ağırlık (kg), Hacim (m3)")
    p.add_argument("--araclar", type=Path, default=ORNEK / "araclar.csv", help="Araç Tipi, Palet Kapasitesi, Ağırlık Kapasitesi (kg), Hacim Kapasitesi (m3), Sefer Maliyeti, Mevcut Adet")
    p.add_argument("--min-doluluk", type=Decimal, default=Decimal(60), help="Bu oranın altındaki seferler için uyarı (varsayılan %%60)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "arac_plani.xlsx")
    a = p.parse_args(argv)
    for y in (a.siparisler, a.araclar):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.siparisler, a.araclar, a.cikti, a.min_doluluk)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    sf = r["seferler"]
    ort = yuvarla(sum(s.doluluk()[0] for s in sf) / len(sf), "0.1") if sf else 0
    print(f"[OK] {r['siparis']} sipariş · {len(sf)} sefer · ortalama doluluk %{ort} · toplam maliyet {sum(s.arac.maliyet for s in sf):,.0f} TL".replace(",", "."))
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['kim']} · {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
