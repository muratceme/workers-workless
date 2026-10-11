import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

G1, G2 = date(2026, 10, 12), date(2026, 10, 13)


def araclar(tir_adet=None):
    return [main.Arac("Panelvan", D(3), D(1000), D(10), D(6500)), main.Arac("Kamyonet", D(8), D(3000), D(20), D(9500)),
            main.Arac("Kamyon", D(18), D(13000), D(45), D(16000)), main.Arac("TIR", D(33), D(24000), D(86), D(24000), tir_adet)]


def sip(no, palet, kg, m3=0, tarih=G1, bolge="İzmir"):
    return main.Siparis(no, "M", bolge, tarih, D(str(palet)), D(str(kg)), D(str(m3)))


def ozet(r):
    return [(s.arac.tip, sorted(x.no for x in s.siparisler)) for s in r["seferler"]]


class PlanTesti(unittest.TestCase):
    def test_tek_siparis_ve_doluluk(self):
        r = main.planla([sip("A", 5, 2000, 9)], araclar())
        self.assertEqual(ozet(r), [("Kamyonet", ["A"])])
        d, kisit = r["seferler"][0].doluluk()
        self.assertEqual((main.yuvarla(d, "0.1"), kisit), (D("66.7"), "Ağırlık"))          # 5/8 = %62,5; 2000/3000 = %66,7; 9/20 = %45

    def test_kucultme(self):
        self.assertEqual(ozet(main.planla([sip("A", 1, 300, 1)], araclar())), [("Panelvan", ["A"])])

    def test_tek_tip_daha_ucuzsa(self):
        # karma sezgisel iki kamyonet (19.000) seçerdi; tek kamyon 16.000
        r = main.planla([sip("A", 4, 2760), sip("B", 2, 1560)], araclar())
        self.assertEqual(ozet(r), [("Kamyon", ["A", "B"])])

    def test_gruplar_birlesmez(self):
        r = main.planla([sip("A", 2, 800), sip("B", 2, 800, tarih=G2), sip("C", 2, 800, bolge="Ankara")], araclar())
        self.assertEqual(len(r["seferler"]), 3)
        self.assertEqual({s.arac.tip for s in r["seferler"]}, {"Panelvan"})

    def test_bolme(self):
        r = main.planla([sip("A", 40, 26000, 72)], araclar())
        self.assertEqual([(s.arac.tip, s.yuk("palet"), s.yuk("kg"), s.yuk("m3")) for s in r["seferler"]],
                         [("TIR", D(33), D("21450.00"), D("59.400")), ("Kamyon", D(7), D("4550.00"), D("12.600"))])   # oran 33/40
        self.assertIn("Sipariş bölündü", [u["tur"] for u in r["uyarilar"]])
        self.assertEqual([x.parca for s in r["seferler"] for x in s.siparisler], ["1/2", "2/2"])

    def test_agirlik_siniri(self):
        r = main.planla([sip("A", 10, 12000), sip("B", 10, 12000), sip("C", 10, 12000)], araclar())    # palet sığar, ağırlık sığmaz
        self.assertEqual(sorted(s.arac.tip for s in r["seferler"]), ["Kamyon", "TIR"])                  # TIR iki sipariş (24.000 kg) + kamyon: 40.000
        self.assertEqual(sum(s.arac.maliyet for s in r["seferler"]), D(40000))                          # 3 kamyon veya 2 TIR 48.000 olurdu
        self.assertTrue(all(s.yuk("kg") <= s.arac.kg for s in r["seferler"]))

    def test_arac_adedi(self):
        r = main.planla([sip("A", 30, 9000), sip("B", 30, 9000)], araclar(tir_adet=1))
        self.assertEqual(ozet(r), [("TIR", ["A"])])
        self.assertEqual([u["tur"] for u in r["uyarilar"] if u["onem"] == "Yüksek"], ["Araç yetersiz"])
        r = main.planla([sip("A", 30, 9000), sip("B", 30, 9000, tarih=G2)], araclar(tir_adet=1))        # adet günlük sayılır
        self.assertEqual(len(r["seferler"]), 2)

    def test_dusuk_doluluk(self):
        r = main.planla([sip("A", 10, 3500), sip("B", 2, 500, tarih=G2)], araclar())
        u = [x for x in r["uyarilar"] if x["tur"] == "Düşük doluluk"]
        self.assertEqual(len(u), 1)                                                        # kamyon 10/18 = %55,6; panelvan %66,7
        self.assertIn("bir gün arayla", u[0]["aciklama"])
        self.assertEqual(len([x for x in main.planla([sip("A", 10, 3500)], araclar(), D(50))["uyarilar"]]), 0)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "p.xlsx"
            r = main.calistir(main.ORNEK / "siparisler.csv", main.ORNEK / "araclar.csv", cikti)
            self.assertEqual(r["siparis"], 34)
            yuklenen = sum(x.palet for s in r["seferler"] for x in s.siparisler)
            self.assertEqual(yuklenen, D(171))                                              # tüm paletler bir araçta
            for s in r["seferler"]:
                self.assertTrue(all(s.yuk(b) <= getattr(s.arac, b) for b in ("palet", "kg", "m3")))
            self.assertEqual({u["tur"] for u in r["uyarilar"]}, {"Teslim tarihi okunamadı", "Düşük doluluk", "Sipariş bölündü"})
            self.assertEqual(load_workbook(cikti).sheetnames, ["Araç Planı", "Yükleme Listesi", "Özet", "Uyarılar"])

    def test_okuma_hatalari(self):
        with tempfile.TemporaryDirectory() as tmp:
            s, a = Path(tmp) / "s.csv", Path(tmp) / "a.csv"
            s.write_text("Sipariş No;Teslim Tarihi;Palet;Ağırlık (kg)\nA;12.10.2026;2;\nB;12.10.2026;0;0\nC;bilinmiyor;1;100\n", encoding="utf-8")
            a.write_text("Araç Tipi;Palet Kapasitesi;Sefer Maliyeti\nKamyonet;8;9.500,00\nBozuk;;1000\n", encoding="utf-8")
            sipler, arac, uy = main.oku(s, a)
            self.assertEqual(([x.no for x in sipler], [x.tip for x in arac], arac[0].maliyet), (["A"], ["Kamyonet"], D("9500.00")))
            self.assertEqual(sorted(u["tur"] for u in uy), ["Araç tanımı eksik", "Eksik boyut", "Teslim tarihi okunamadı", "Yük bilgisi yok"])
            self.assertEqual(main.main(["--siparisler", str(s), "--araclar", str(a), "--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--siparisler", str(Path(tmp) / "yok.csv")]), 1)
            a.write_text("Araç Tipi;Sefer Maliyeti\nX;0\n", encoding="utf-8")
            self.assertEqual(main.main(["--siparisler", str(s), "--araclar", str(a)]), 1)


if __name__ == "__main__":
    unittest.main()
