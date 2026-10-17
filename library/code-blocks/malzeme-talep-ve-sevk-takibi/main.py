"""
Malzeme Talep ve Sevk Takibi — Gerisi Bende kod bloğu
Şantiye › Şantiye Şefi
Şantiye malzeme taleplerini satın alma ve sevkiyat durumuyla takip eder, imalatı geciktirecek eksikleri işaretler.
  Durum            = Teslim edildi (teslim ≥ talep) · Kısmi teslim · Siparişte (sipariş var, teslim yok) · Sipariş bekliyor (sipariş yok)
  Kalan miktar     = talep miktarı − teslim edilen
  Riskler          = İhtiyaç tarihi geçti, malzeme tam gelmedi (Yüksek) · Termin ihtiyaç tarihinden sonra (Orta) ·
                     Sipariş verilmemiş ve ihtiyaca --tedarik (7) gün ya da daha az kalmış (Orta)
  Tedarikçi        = zamanında teslim % (son teslim ≤ termin) ve ortalama talep → teslim süresi
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"no": ("talep no", "talep", "talep numarasi"), "tarih": ("talep tarihi", "tarih"), "malzeme": ("malzeme", "malzeme adi", "kalem", "urun"), "miktar": ("miktar", "talep miktari", "istenen miktar"),
         "birim": ("birim", "olcu birimi"), "ihtiyac": ("ihtiyac tarihi", "gerekli tarih", "istenen tarih"), "imalat": ("imalat", "imalat yer", "yer", "blok", "kullanim yeri", "is kalemi"),
         "siparis": ("siparis no", "siparis", "po no"), "siparis_tarih": ("siparis tarihi",), "tedarikci": ("tedarikci", "firma", "satici"), "termin": ("termin", "teslim termini", "soz verilen tarih"),
         "teslim": ("teslim edilen", "teslim miktari", "gelen miktar", "sevk edilen"), "teslim_tarih": ("teslim tarihi", "son teslim tarihi", "sevk tarihi", "gelis tarihi")}
SIFIR = Decimal(0)
DURUM_SIRA = {"Sipariş bekliyor": 0, "Siparişte": 1, "Kısmi teslim": 2, "Teslim edildi": 3}


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("malzeme", "miktar", "ihtiyac")):
        malzeme = metin(r["malzeme"])
        if not malzeme:
            continue
        no = metin(r["no"])
        tarihler = {a: tarih(r[a]) if metin(r[a]) else None for a in ("tarih", "ihtiyac", "siparis_tarih", "termin", "teslim_tarih")}
        miktar, teslim = para(r["miktar"]), para(r["teslim"]) if metin(r["teslim"]) else SIFIR
        kim = f"{no or r['_satir']} · {malzeme}"
        if miktar is None or miktar <= 0 or teslim is None or teslim < 0 or tarihler["ihtiyac"] is None or any(metin(r[a]) and v is None for a, v in tarihler.items()):
            uy.append(uyari("Yüksek", "Satır okunamadı", kim, f"Satır {r['_satir']}: miktar, teslim miktarı ya da tarihlerden biri okunamadı. Satır alınmadı."))
            continue
        if no and (katla(no), katla(malzeme)) in gorulen:
            uy.append(uyari("Orta", "Talep kalemi iki kez yazılmış", kim, f"Satır {r['_satir']}: ikinci satır alınmadı."))
            continue
        gorulen.add((katla(no), katla(malzeme)))
        l.append({"no": no, "malzeme": malzeme, "miktar": miktar, "birim": metin(r["birim"]), "imalat": metin(r["imalat"]) or "-", "siparis": metin(r["siparis"]), "tedarikci": metin(r["tedarikci"]),
                  "teslim": teslim, **tarihler})
    return l, uy


def takip(talepler: list[dict], gun: date, tedarik: int = 7) -> dict:
    if not talepler:
        raise ValueError("Talep kaydı yok.")
    uy, sonuc = [], []
    for t in talepler:
        kalan = max(t["miktar"] - t["teslim"], SIFIR)
        durum = "Teslim edildi" if t["teslim"] >= t["miktar"] else "Kısmi teslim" if t["teslim"] > 0 else "Siparişte" if t["siparis"] or t["siparis_tarih"] else "Sipariş bekliyor"
        risk, kim = [], f"{t['no'] or '-'} · {t['malzeme']}"
        acik = durum != "Teslim edildi"
        if acik and t["ihtiyac"] < gun:
            risk.append("İhtiyaç tarihi geçti")
            uy.append(uyari("Yüksek", "İhtiyaç tarihi geçti", kim, f"{t['imalat']}: ihtiyaç {t['ihtiyac']:%d.%m.%Y}, {(gun - t['ihtiyac']).days} gün geçti; kalan {kalan:,.2f} {t['birim']} ({durum.lower()})."))
        if acik and t["termin"] and t["termin"] > t["ihtiyac"]:
            risk.append("Termin ihtiyaçtan sonra")
            uy.append(uyari("Orta", "Termin ihtiyaçtan sonra", kim, f"{t['imalat']}: termin {t['termin']:%d.%m.%Y}, ihtiyaç {t['ihtiyac']:%d.%m.%Y}" + (f" ({t['tedarikci']})." if t["tedarikci"] else ".")))
        if durum == "Sipariş bekliyor" and t["ihtiyac"] >= gun and (t["ihtiyac"] - gun).days <= tedarik:
            risk.append("Sipariş verilmedi")
            uy.append(uyari("Orta", "Sipariş verilmedi, süre daralıyor", kim, f"{t['imalat']}: ihtiyaca {(t['ihtiyac'] - gun).days} gün kaldı (tedarik süresi {tedarik} gün)."))
        if not acik and t["teslim_tarih"] and t["teslim_tarih"] > t["ihtiyac"]:
            uy.append(uyari("Bilgi", "Geç teslim", kim, f"Teslim {t['teslim_tarih']:%d.%m.%Y}, ihtiyaç {t['ihtiyac']:%d.%m.%Y}."))
        if t["teslim"] > t["miktar"]:
            uy.append(uyari("Bilgi", "Talepten fazla teslim", kim, f"Talep {t['miktar']:,.2f}, teslim {t['teslim']:,.2f} {t['birim']}."))
        sonuc.append({**t, "durum": durum, "kalan": kalan, "risk": ", ".join(risk), "kalan_gun": (t["ihtiyac"] - gun).days})
    sonuc.sort(key=lambda s: (not s["risk"], DURUM_SIRA[s["durum"]] if s["risk"] else 9, s["ihtiyac"], katla(s["malzeme"])))
    ted = {}
    for s in sonuc:
        if s["tedarikci"] and s["teslim_tarih"]:
            x = ted.setdefault(katla(s["tedarikci"]), {"ad": s["tedarikci"], "teslim": 0, "zamaninda": 0, "terminli": 0, "sureler": []})
            x["teslim"] += 1
            if s["termin"]:
                x["terminli"] += 1
                x["zamaninda"] += s["teslim_tarih"] <= s["termin"]
            if s["tarih"]:
                x["sureler"].append((s["teslim_tarih"] - s["tarih"]).days)
    tedarikci = sorted(({"ad": x["ad"], "teslim": x["teslim"], "oran": yuvarla(Decimal(x["zamaninda"]) * 100 / x["terminli"]) if x["terminli"] else None,
                         "sure": yuvarla(Decimal(sum(x["sureler"])) / len(x["sureler"])) if x["sureler"] else None} for x in ted.values()), key=lambda x: katla(x["ad"]))
    ozet = {d: sum(1 for s in sonuc if s["durum"] == d) for d in DURUM_SIRA}
    return {"talepler": sonuc, "tedarikci": tedarikci, "ozet": ozet, "riskli": sum(1 for s in sonuc if s["risk"]), "uyarilar": uy}


DURUM_RENK = {"Sipariş bekliyor": SARI, "Siparişte": MAVI, "Kısmi teslim": SARI, "Teslim edildi": YESIL}


def rapor_yaz(cikti: Path, r: dict, gun: date) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Talepler"
    _baslik(ws, ["Risk", "Durum", "Talep No", "Talep Tarihi", "Malzeme", "Miktar", "Birim", "Teslim Edilen", "Kalan", "İmalat / Yer", "İhtiyaç Tarihi", "Kalan Gün", "Sipariş No", "Tedarikçi",
                 "Termin", "Teslim Tarihi"], (26, 14, 10, 11, 26, 9, 7, 9, 9, 22, 11, 8, 11, 18, 11, 11), "E2")
    for s in r["talepler"]:
        _satir(ws, [s["risk"], s["durum"], s["no"], s["tarih"], s["malzeme"], s["miktar"], s["birim"], s["teslim"] or None, s["kalan"] or None, s["imalat"], s["ihtiyac"],
                    s["kalan_gun"] if s["durum"] != "Teslim edildi" else None, s["siparis"], s["tedarikci"], s["termin"], s["teslim_tarih"]], tarih_s=(4, 11, 15, 16),
               dolgu={1: KIRMIZI if "geçti" in s["risk"] else SARI if s["risk"] else None, 2: DURUM_RENK[s["durum"]]})
    ws.append([])
    ws.append([f"Rapor tarihi {gun:%d.%m.%Y}. Riskli kalemler üstte."])
    ws = wb.create_sheet("Tedarikçiler")
    _baslik(ws, ["Tedarikçi", "Teslim Edilen Kalem", "Zamanında Teslim %", "Ortalama Talep → Teslim (gün)"], (24, 10, 12, 14))
    for x in r["tedarikci"]:
        _satir(ws, [x["ad"], x["teslim"], x["oran"], x["sure"]], para_s=(3, 4))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, gun: date, tedarik: int = 7) -> dict:
    talepler, uy = oku(yol)
    r = takip(talepler, gun, tedarik)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, gun)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Şantiye malzeme taleplerini satın alma ve sevkiyat durumuyla takip eder.")
    p.add_argument("--talepler", type=Path, help="Talep No, Talep Tarihi, Malzeme, Miktar, Birim, İhtiyaç Tarihi, İmalat, Sipariş No, Sipariş Tarihi, Tedarikçi, Termin, Teslim Edilen, Teslim Tarihi")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 15.10.2026)")
    p.add_argument("--tedarik", type=int, default=7, help="Sipariş verilmemiş kalemde ihtiyaca bu kadar gün kalınca uyarılır (varsayılan 7)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "malzeme_talep_takibi.xlsx")
    a = p.parse_args(argv)
    yol = a.talepler or ORNEK / "malzeme_talepleri.csv"
    gun = tarih(a.tarih) if a.tarih else (date(2026, 10, 15) if a.talepler is None else date.today())
    if gun is None or a.tedarik < 0:
        print("[X] --tarih GG.AA.YYYY, --tedarik sıfır ya da üstü olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, gun, a.tedarik)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {len(r['talepler'])} kalem · sipariş bekliyor {o['Sipariş bekliyor']} · siparişte {o['Siparişte']} · kısmi {o['Kısmi teslim']} · teslim edildi {o['Teslim edildi']} · riskli {r['riskli']}")
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
