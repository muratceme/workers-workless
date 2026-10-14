import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
GUN = date(2026, 10, 15)


def t(no, talep, malzeme, miktar, ihtiyac, siparis="", ted="", termin=None, teslim=0, teslim_t=None):
    g = lambda x: date(2026, 10, x) if x else None   # noqa: E731
    return {"no": no, "malzeme": malzeme, "miktar": D(miktar), "birim": "ad", "imalat": "A Blok", "siparis": siparis, "tedarikci": ted, "teslim": D(teslim), "tarih": g(talep),
            "ihtiyac": g(ihtiyac), "siparis_tarih": None, "termin": g(termin), "teslim_tarih": g(teslim_t)}


def veri():
    return [t("T1", 1, "M1", 100, 10, "S1", "A", 8, 100, 7), t("T2", 1, "M2", 50, 12, "S2", "A", 14, 20, 12), t("T3", 5, "M3", 10, 20), t("T4", 5, "M4", 5, 30),
            t("T5", 1, "M5", 30, 9, "S3", "B", 8, 35, 10)]


class TakipTesti(unittest.TestCase):
    def test_durum_ve_risk(self):
        r = main.takip(veri(), GUN)
        self.assertEqual([(s["no"], s["durum"], s["kalan"], s["risk"], s["kalan_gun"]) for s in r["talepler"]],
                         [("T3", "Sipariş bekliyor", D(10), "Sipariş verilmedi", 5), ("T2", "Kısmi teslim", D(30), "İhtiyaç tarihi geçti, Termin ihtiyaçtan sonra", -3),
                          ("T5", "Teslim edildi", D(0), "", -6), ("T1", "Teslim edildi", D(0), "", -5), ("T4", "Sipariş bekliyor", D(5), "", 15)])
        self.assertEqual([(x["ad"], x["teslim"], x["oran"], x["sure"]) for x in r["tedarikci"]], [("A", 2, D("100.00"), D("8.50")), ("B", 1, D("0.00"), D("9.00"))])
        self.assertEqual((r["ozet"], r["riskli"]), ({"Sipariş bekliyor": 2, "Siparişte": 0, "Kısmi teslim": 1, "Teslim edildi": 2}, 2))
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Bilgi", "Geç teslim", "T5 · M5"), ("Bilgi", "Talepten fazla teslim", "T5 · M5"), ("Orta", "Sipariş verilmedi, süre daralıyor", "T3 · M3"),
                          ("Orta", "Termin ihtiyaçtan sonra", "T2 · M2"), ("Yüksek", "İhtiyaç tarihi geçti", "T2 · M2")])

    def test_tedarik(self):
        r = main.takip(veri(), GUN, tedarik=15)
        self.assertEqual(sorted(u["kim"] for u in r["uyarilar"] if u["tur"].startswith("Sipariş verilmedi")), ["T3 · M3", "T4 · M4"])
        with self.assertRaises(ValueError):
            main.takip([], GUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "malzeme_talepleri.csv", cikti, GUN)
            self.assertEqual((r["ozet"], r["riskli"]), ({"Sipariş bekliyor": 2, "Siparişte": 3, "Kısmi teslim": 1, "Teslim edildi": 4}, 3))
            self.assertEqual([(x["ad"], x["oran"]) for x in r["tedarikci"]], [("Örnek Demir Ticaret", D("50.00")), ("Örnek Hazır Beton", D("100.00")), ("Örnek Kalıp Malzemeleri", D("50.00"))])
            self.assertEqual([u["tur"] for u in r["uyarilar"]], ["İhtiyaç tarihi geçti", "İhtiyaç tarihi geçti", "Sipariş verilmedi, süre daralıyor", "Termin ihtiyaçtan sonra", "Geç teslim", "Geç teslim"])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Talepler", "Tedarikçiler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "t.csv"
            y.write_text("Talep No;Kalem;Miktar;Gerekli Tarih;Gelen Miktar\nT1;Çimento;1.200;20.10.2026;\nT1;çimento;5;20.10.2026;\nT2;Kum;x;20.10.2026;\nT3;Kum;5;;\nT4;Kum;5;20.10.2026;-1\n;Taş;5;20.10.2026;2\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["miktar"], x["teslim"]) for x in l], [u["tur"] for u in uy]),
                             ([("T1", D(1200), D(0)), ("", D(5), D(2))], ["Talep kalemi iki kez yazılmış", "Satır okunamadı", "Satır okunamadı", "Satır okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--talepler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "15.10.2026", "--tedarik", "10"]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--tedarik", "-1"])], [2] * 2)
            self.assertEqual(main.main(["--talepler", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Talep No;Kalem\nT1;x\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
