"""
Sürücü Çalışma ve Dinlenme Süresi Kontrolü — Gerisi Bende kod bloğu
Filo Yönetimi › Filo Yöneticisi

Takograf verisinden günlük ve haftalık sürüş ile dinlenme sürelerini AETR kurallarına göre kontrol eder, ihlalleri listeler.
Girdi günlük özet satırlarıdır (sürücü, tarih, günlük sürüş, en uzun kesintisiz sürüş, günlük dinlenme). Kontroller ve varsayılan sınırlar:
  Günlük sürüş      : 9 saati aşan gün "uzatılmış" sayılır; haftada en çok 2 uzatılmış güne izin vardır; 10 saatin üzeri her durumda ihlaldir
  Kesintisiz sürüş  : en çok 4,5 saat (sonrasında en az 45 dakika mola gerekir; molanın süresi ve bölünmesi bu pakette kontrol edilmez)
  Günlük dinlenme   : en az 11 saat; 9 saatten az ise ihlal; 9-11 saat "azaltılmış" sayılır ve haftada en çok 3 kez olabilir
  Haftalık sürüş    : pazartesi-pazar toplamı en çok 56 saat
  İki haftalık sürüş: art arda iki haftanın toplamı en çok 90 saat
Sınırlar parametreyle değişir. Bu pakette kaynak teyidi yapılmamıştır: güncel AETR metnini ve Karayolları Trafik mevzuatını siz doğrulayın.
Kapsam dışı: haftalık dinlenme (45 / 24 saat) ve telafisi, bölünmüş günlük dinlenme, ekipli (çift sürücülü) sürüş, feribot / tren istisnaları, diğer çalışma süreleri.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"surucu": ("surucu", "sofor", "surucu adi", "kart no"), "tarih": ("tarih", "gun"), "surus": ("surus suresi", "gunluk surus", "surus", "toplam surus"),
         "kesintisiz": ("en uzun kesintisiz surus", "kesintisiz surus", "kesintisiz", "araliksiz surus"), "dinlenme": ("gunluk dinlenme", "dinlenme", "dinlenme suresi"), "plaka": ("plaka", "arac")}
VARSAYILAN = {"gunluk": Decimal(9), "gunluk_uzatilmis": Decimal(10), "uzatma_hakki": 2, "kesintisiz": Decimal("4.5"), "dinlenme": Decimal(11), "dinlenme_azaltilmis": Decimal(9), "azaltma_hakki": 3,
              "haftalik": Decimal(56), "iki_haftalik": Decimal(90)}
Z = Decimal(0)


def sure(x) -> Decimal | None:
    """9 · 9,5 · 9.5 · 9:30 · 09:30:00 · saat nesnesi → ondalık saat"""
    if x in (None, ""):
        return None
    if hasattr(x, "hour") and not isinstance(x, (date, datetime)) or isinstance(x, datetime):
        return Decimal(x.hour) + Decimal(x.minute) / 60
    if isinstance(x, timedelta):
        return Decimal(int(x.total_seconds())) / 3600
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = metin(x)
    m = re.fullmatch(r"(\d{1,2}):(\d{2})(?::\d{2})?", s)
    if m and int(m[2]) < 60:
        return Decimal(int(m[1])) + Decimal(int(m[2])) / 60
    if re.fullmatch(r"\d{1,2}([.,]\d{1,2})?", s):
        return Decimal(s.replace(",", "."))
    return None


def saat_yaz(s: Decimal) -> str:
    dk = int((s * 60).to_integral_value(rounding=ROUND_HALF_UP))
    return f"{dk // 60}:{dk % 60:02d}"


def pazartesi(d: date) -> date:
    return d - timedelta(days=d.weekday())


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("surucu", "tarih", "surus")):
        s = metin(r["surucu"])
        if not s:
            continue
        kim = f"{s} (satır {r['_satir']})"
        t, surus = tarih(r["tarih"]), sure(r["surus"])
        if t is None or surus is None or surus > 24:
            uy.append(uyari("Yüksek", "Kayıt okunamadı", kim, f"Tarih '{metin(r['tarih'])}', sürüş süresi '{metin(r['surus'])}'. Süre saat (9,5) veya saat:dakika (9:30) olmalı. Satır alınmadı."))
            continue
        if (katla(s), t) in gorulen:
            uy.append(uyari("Yüksek", "Mükerrer gün", kim, f"{t:%d.%m.%Y} için birden fazla satır; yalnız ilki alındı. Günlük toplamı tek satırda verin."))
            continue
        gorulen.add((katla(s), t))
        x = {"surucu": s, "tarih": t, "surus": surus, "plaka": metin(r["plaka"])}
        for a, ad in (("kesintisiz", "Kesintisiz sürüş"), ("dinlenme", "Günlük dinlenme")):
            x[a] = sure(r[a])
            if x[a] is None and metin(r[a]) or (x[a] is not None and x[a] > 24):
                uy.append(uyari("Orta", "Süre okunamadı", kim, f"{ad} '{metin(r[a])}'. Bu kontrol o gün için yapılmadı."))
                x[a] = None
        if x["kesintisiz"] is not None and x["kesintisiz"] > surus:
            uy.append(uyari("Orta", "Kesintisiz sürüş günlük sürüşten uzun", kim, f"Kesintisiz {saat_yaz(x['kesintisiz'])}, günlük {saat_yaz(surus)}. Sütunları kontrol edin."))
        l.append(x)
    return l, uy


def kontrol(kayit: list[dict], sinir: dict | None = None) -> dict:
    if not kayit:
        raise ValueError("Kontrol edilecek kayıt yok.")
    s, ihlal, haftalar, ozet = VARSAYILAN | (sinir or {}), [], [], []

    def ekle(k, donem, kural, deger, limit, aciklama, tarih_=None):
        ihlal.append({"surucu": k["surucu"], "plaka": k.get("plaka", ""), "tarih": tarih_, "donem": donem, "kural": kural, "deger": deger, "sinir": limit, "asim": abs(deger - limit), "aciklama": aciklama})

    grup = defaultdict(list)
    for k in kayit:
        grup[katla(k["surucu"])].append(k)
    for liste in grup.values():
        liste.sort(key=lambda k: k["tarih"])
        n0, hafta = len(ihlal), defaultdict(list)
        for k in liste:
            hafta[pazartesi(k["tarih"])].append(k)
        toplam = {}
        for h, gunler in sorted(hafta.items()):
            uzatma = azaltma = 0
            for k in gunler:
                gun = f"{k['tarih']:%d.%m.%Y}"
                if k["surus"] > s["gunluk_uzatilmis"]:
                    uzatma += 1
                    ekle(k, gun, "Günlük sürüş", k["surus"], s["gunluk_uzatilmis"], f"Günlük sürüş {saat_yaz(k['surus'])}; uzatılmış sınır {saat_yaz(s['gunluk_uzatilmis'])} saat.", k["tarih"])
                elif k["surus"] > s["gunluk"]:
                    uzatma += 1
                    if uzatma > s["uzatma_hakki"]:
                        ekle(k, gun, "Uzatılmış sürüş hakkı", k["surus"], s["gunluk"], f"Bu hafta {uzatma}. kez {saat_yaz(s['gunluk'])} saat aşıldı; haftada en çok {s['uzatma_hakki']} kez uzatılabilir.", k["tarih"])
                if k["kesintisiz"] is not None and k["kesintisiz"] > s["kesintisiz"]:
                    ekle(k, gun, "Kesintisiz sürüş", k["kesintisiz"], s["kesintisiz"], f"Mola vermeden {saat_yaz(k['kesintisiz'])} sürüş; sınır {saat_yaz(s['kesintisiz'])} saat.", k["tarih"])
                if k["dinlenme"] is not None:
                    if k["dinlenme"] < s["dinlenme_azaltilmis"]:
                        ekle(k, gun, "Günlük dinlenme", k["dinlenme"], s["dinlenme_azaltilmis"], f"Günlük dinlenme {saat_yaz(k['dinlenme'])}; azaltılmış dinlenme bile en az {saat_yaz(s['dinlenme_azaltilmis'])} saat.", k["tarih"])
                    elif k["dinlenme"] < s["dinlenme"]:
                        azaltma += 1
                        if azaltma > s["azaltma_hakki"]:
                            ekle(k, gun, "Azaltılmış dinlenme hakkı", k["dinlenme"], s["dinlenme"], f"Bu hafta {azaltma}. kez {saat_yaz(s['dinlenme'])} saatin altında dinlenme; en çok {s['azaltma_hakki']} kez azaltılabilir.", k["tarih"])
            toplam[h] = sum((k["surus"] for k in gunler), Z)
            etiket = f"{h:%d.%m.%Y} haftası"
            if toplam[h] > s["haftalik"]:
                ekle(gunler[0], etiket, "Haftalık sürüş", toplam[h], s["haftalik"], f"Haftalık toplam {saat_yaz(toplam[h])}; sınır {saat_yaz(s['haftalik'])} saat.")
            onceki = toplam.get(h - timedelta(days=7))
            iki = toplam[h] + onceki if onceki is not None else None
            if iki is not None and iki > s["iki_haftalik"]:
                ekle(gunler[0], f"{h - timedelta(days=7):%d.%m.%Y} ve {h:%d.%m.%Y} haftaları", "İki haftalık sürüş", iki, s["iki_haftalik"], f"İki haftanın toplamı {saat_yaz(iki)}; sınır {saat_yaz(s['iki_haftalik'])} saat.")
            haftalar.append({"surucu": liste[0]["surucu"], "hafta": h, "gun": len(gunler), "surus": toplam[h], "iki": iki, "uzatma": uzatma, "azaltma": azaltma})
        kendi = ihlal[n0:]
        ozet.append({"surucu": liste[0]["surucu"], "gun": len(liste), "surus": sum((k["surus"] for k in liste), Z), "ihlal": len(kendi), "kurallar": sorted({i["kural"] for i in kendi}),
                     "eksik": sum(k["kesintisiz"] is None or k["dinlenme"] is None for k in liste)})
    return {"ihlaller": sorted(ihlal, key=lambda i: (i["surucu"], i["tarih"] or date.max, i["kural"])), "ozet": sorted(ozet, key=lambda o: (-o["ihlal"], o["surucu"])), "haftalar": haftalar,
            "kayitlar": sorted(kayit, key=lambda k: (k["surucu"], k["tarih"])), "sinir": s, "say": {"kayit": len(kayit), "ihlal": len(ihlal), "ihlalli_surucu": sum(o["ihlal"] > 0 for o in ozet)}}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    s = r["sinir"]
    ws = wb.active
    ws.title = "İhlaller"
    _baslik(ws, ["Sürücü", "Plaka", "Gün / Dönem", "Kural", "Gerçekleşen (saat)", "Sınır (saat)", "Fark (saat)", "Açıklama", "Alınan Önlem (siz doldurun)"], (24, 12, 30, 24, 11, 9, 9, 70, 30))
    for i in r["ihlaller"]:
        _satir(ws, [i["surucu"], i["plaka"], i["donem"], i["kural"], saat_yaz(i["deger"]), saat_yaz(i["sinir"]), saat_yaz(i["asim"]), i["aciklama"], ""], dolgu={4: KIRMIZI}, sar=(8,))
    ws.append([])
    ws.append([f"Sınırlar: günlük {saat_yaz(s['gunluk'])} (haftada {s['uzatma_hakki']} kez {saat_yaz(s['gunluk_uzatilmis'])}), kesintisiz {saat_yaz(s['kesintisiz'])}, günlük dinlenme {saat_yaz(s['dinlenme'])} "
               f"(haftada {s['azaltma_hakki']} kez en az {saat_yaz(s['dinlenme_azaltilmis'])}), haftalık {saat_yaz(s['haftalik'])}, iki haftalık {saat_yaz(s['iki_haftalik'])} saat. "
               "Haftalık dinlenme ve mola süreleri bu pakette kontrol edilmez. Güncel mevzuatı doğrulayın."])
    ws = wb.create_sheet("Sürücü Özeti")
    _baslik(ws, ["Sürücü", "Kayıtlı Gün", "Toplam Sürüş (saat)", "İhlal Sayısı", "İhlal Edilen Kurallar", "Verisi Eksik Gün"], (26, 9, 12, 9, 60, 10))
    for o in r["ozet"]:
        _satir(ws, [o["surucu"], o["gun"], saat_yaz(o["surus"]), o["ihlal"] or None, ", ".join(o["kurallar"]), o["eksik"] or None], dolgu={4: KIRMIZI if o["ihlal"] else YESIL}, sar=(5,))
    ws = wb.create_sheet("Haftalık Toplamlar")
    _baslik(ws, ["Sürücü", "Hafta (pazartesi)", "Kayıtlı Gün", "Haftalık Sürüş (saat)", "Önceki Haftayla Toplam", "9 Saati Aşan Gün", "Azaltılmış Dinlenme"], (26, 14, 9, 12, 12, 10, 10))
    for h in r["haftalar"]:
        _satir(ws, [h["surucu"], h["hafta"], h["gun"], saat_yaz(h["surus"]), saat_yaz(h["iki"]) if h["iki"] is not None else None, h["uzatma"] or None, h["azaltma"] or None], tarih_s=(2,),
               dolgu={4: KIRMIZI if h["surus"] > s["haftalik"] else None, 5: KIRMIZI if h["iki"] is not None and h["iki"] > s["iki_haftalik"] else None})
    ws = wb.create_sheet("Günlük Kayıtlar")
    _baslik(ws, ["Sürücü", "Tarih", "Plaka", "Günlük Sürüş", "En Uzun Kesintisiz Sürüş", "Günlük Dinlenme"], (26, 12, 12, 11, 12, 11))
    for k in r["kayitlar"]:
        _satir(ws, [k["surucu"], k["tarih"], k["plaka"], saat_yaz(k["surus"]), saat_yaz(k["kesintisiz"]) if k["kesintisiz"] is not None else None, saat_yaz(k["dinlenme"]) if k["dinlenme"] is not None else None], tarih_s=(2,))
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, sinir: dict | None = None) -> dict:
    kayit, uy = oku(yol)
    r = kontrol(kayit, sinir)
    for o in r["ozet"]:
        if o["eksik"]:
            uy.append(uyari("Orta", "Eksik veri", o["surucu"], f"{o['eksik']} günde kesintisiz sürüş veya günlük dinlenme boş; bu günlerde ilgili kontrol yapılmadı."))
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Günlük takograf özetlerinden sürüş ve dinlenme süresi ihlallerini listeler.")
    p.add_argument("--kayitlar", type=Path, default=ORNEK / "takograf_ozeti.csv", help="Sürücü, Tarih, Sürüş Süresi, En Uzun Kesintisiz Sürüş, Günlük Dinlenme, Plaka (süreler saat veya saat:dakika)")
    for ad, anahtar, yardim in (("gunluk-surus", "gunluk", "Günlük sürüş sınırı, saat"), ("uzatilmis-surus", "gunluk_uzatilmis", "Uzatılmış günlük sürüş sınırı"), ("uzatma-hakki", "uzatma_hakki", "Haftada uzatılmış gün sayısı"),
                               ("kesintisiz-surus", "kesintisiz", "Kesintisiz sürüş sınırı"), ("gunluk-dinlenme", "dinlenme", "Günlük dinlenme"), ("azaltilmis-dinlenme", "dinlenme_azaltilmis", "Azaltılmış günlük dinlenme"),
                               ("azaltma-hakki", "azaltma_hakki", "Haftada azaltılmış dinlenme sayısı"), ("haftalik-surus", "haftalik", "Haftalık sürüş sınırı"), ("iki-haftalik-surus", "iki_haftalik", "İki haftalık sürüş sınırı")):
        p.add_argument(f"--{ad}", dest=anahtar, default=str(VARSAYILAN[anahtar]).replace(".", ","), help=f"{yardim} (varsayılan {str(VARSAYILAN[anahtar]).replace('.', ',')})")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "surus_dinlenme_kontrolu.xlsx")
    a = p.parse_args(argv)
    sinir = {k: sure(getattr(a, k)) for k in VARSAYILAN}
    if any(v is None or v < 0 for v in sinir.values()) or sinir["gunluk_uzatilmis"] < sinir["gunluk"] or sinir["dinlenme_azaltilmis"] > sinir["dinlenme"] or any(sinir[k] != int(sinir[k]) for k in ("uzatma_hakki", "azaltma_hakki")):
        print("[X] Sınırlar sıfır veya pozitif olmalı; uzatılmış sürüş günlük sürüşten küçük, azaltılmış dinlenme günlük dinlenmeden büyük olamaz; hak sayıları tam sayı olmalı")
        return 2
    sinir["uzatma_hakki"], sinir["azaltma_hakki"] = int(sinir["uzatma_hakki"]), int(sinir["azaltma_hakki"])
    if not a.kayitlar.exists():
        print(f"[X] Dosya bulunamadı: {a.kayitlar}")
        return 1
    try:
        r = calistir(a.kayitlar, a.cikti, sinir)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    y = r["say"]
    print(f"[OK] {len(r['ozet'])} sürücü · {y['kayit']} günlük kayıt · {y['ihlal']} ihlal ({y['ihlalli_surucu']} sürücüde)")
    for i in r["ihlaller"]:
        print(f"[!] {i['surucu']} · {i['donem']} · {i['kural']}: {saat_yaz(i['deger'])} (sınır {saat_yaz(i['sinir'])})")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
