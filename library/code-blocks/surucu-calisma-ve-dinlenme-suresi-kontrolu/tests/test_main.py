import sys
import tempfile
import unittest
from datetime import date, time, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
PZT = date(2026, 10, 5)


def gn(gun, surus, kesintisiz="4.5", dinlenme=11, surucu="S"):
    f = lambda v: None if v is None else D(v)   # noqa: E731
    return {"surucu": surucu, "tarih": PZT + timedelta(days=gun), "surus": D(surus), "kesintisiz": f(kesintisiz), "dinlenme": f(dinlenme), "plaka": ""}


def kurallar(*gunler, **sinir):
    return [(i["tarih"].day if i["tarih"] else None, i["kural"]) for i in main.kontrol(list(gunler), sinir or None)["ihlaller"]]


class KontrolTesti(unittest.TestCase):
    def test_sure_cozme(self):
        self.assertEqual([main.sure(x) for x in ("9", "9,5", "9.5", "9:30", "09:30:00", time(4, 45), timedelta(hours=10, minutes=15), 8, "dokuz", "9:75", "")],
                         [D(9), D("9.5"), D("9.5"), D("9.5"), D("9.5"), D("4.75"), D("10.25"), D(8), None, None, None])
        self.assertEqual([main.saat_yaz(D(x)) for x in ("4.5", "10.25", "0", "9.99")], ["4:30", "10:15", "0:00", "9:59"])

    def test_gunluk_kurallar(self):
        self.assertEqual(kurallar(gn(0, 9), gn(1, 10), gn(2, 10), gn(3, 9)), [])                                              # haftada iki kez 10 saate kadar
        self.assertEqual(kurallar(gn(0, "9.5"), gn(1, 10), gn(2, "9.25")), [(7, "Uzatılmış sürüş hakkı")])
        self.assertEqual(kurallar(gn(0, "10.25")), [(5, "Günlük sürüş")])
        self.assertEqual(kurallar(gn(0, "10.5"), gn(1, 10), gn(2, 10)), [(5, "Günlük sürüş"), (7, "Uzatılmış sürüş hakkı")])     # 10 saati aşan gün de hakkı kullanır
        self.assertEqual(kurallar(gn(5, 10), gn(6, 10), gn(7, 10), gn(8, 10)), [])                                           # hafta pazartesi yenilenir
        self.assertEqual((kurallar(gn(0, 8, "4.5")), kurallar(gn(0, 8, "4.75")), kurallar(gn(0, 8, None, None))), ([], [(5, "Kesintisiz sürüş")], []))
        self.assertEqual((kurallar(gn(0, 8, dinlenme=9)), kurallar(gn(0, 8, dinlenme="8.75"))), ([], [(5, "Günlük dinlenme")]))
        self.assertEqual(kurallar(*[gn(i, 8, dinlenme=10) for i in range(5)]), [(8, "Azaltılmış dinlenme hakkı"), (9, "Azaltılmış dinlenme hakkı")])
        self.assertEqual(kurallar(gn(0, "9.5"), gunluk=D(10)), [])

    def test_haftalik_kurallar(self):
        self.assertEqual(kurallar(*[gn(i, "9.5" if i < 2 else 9, None, None) for i in range(6)], gn(6, 1, None, None)), [])        # 9,5 + 9,5 + 4 × 9 + 1 = 56
        self.assertEqual(kurallar(*[gn(i, 9, None, None) for i in range(6)], gn(6, "2.25", None, None)), [(None, "Haftalık sürüş")])
        iki = [gn(i, 9, None, None) for i in (0, 1, 2, 3, 4)] + [gn(i, 9, None, None) for i in (7, 8, 9, 10, 11)]
        self.assertEqual((kurallar(*iki), kurallar(*iki, gn(12, "0.25", None, None))), ([], [(None, "İki haftalık sürüş")]))       # 45 + 45 = 90; 90,25 ihlal
        self.assertEqual(kurallar(*[gn(i, 9, None, None) for i in range(6)], *[gn(i, 9, None, None) for i in range(14, 20)]), [])  # arada boş hafta: iki haftalık toplam uygulanmaz
        r = main.kontrol([gn(0, 9), gn(1, 9), gn(0, "10.5", surucu="B")])
        self.assertEqual(([(o["surucu"], o["gun"], o["surus"], o["ihlal"], o["kurallar"]) for o in r["ozet"]], r["say"]), ([("B", 1, D("10.5"), 1, ["Günlük sürüş"]), ("S", 2, D(18), 0, [])], {"kayit": 3, "ihlal": 1, "ihlalli_surucu": 1}))
        with self.assertRaises(ValueError):
            main.kontrol([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "s.xlsx"
            r = main.calistir(ORNEK / "takograf_ozeti.csv", cikti)
            self.assertEqual([(i["surucu"], i["donem"][:10], i["kural"], main.saat_yaz(i["deger"])) for i in r["ihlaller"]],
                             [("Sürücü 1", "06.10.2026", "Günlük dinlenme", "8:00"), ("Sürücü 1", "06.10.2026", "Günlük sürüş", "10:30"), ("Sürücü 1", "06.10.2026", "Kesintisiz sürüş", "5:00"),
                              ("Sürücü 1", "07.10.2026", "Uzatılmış sürüş hakkı", "9:30"), ("Sürücü 1", "08.10.2026", "Uzatılmış sürüş hakkı", "9:30"), ("Sürücü 1", "28.09.2026", "İki haftalık sürüş", "99:00")])
            self.assertEqual(([(h["surucu"], h["surus"], h["iki"]) for h in r["haftalar"]], r["say"], r["uyarilar"]),
                             ([("Sürücü 1", D(52), None), ("Sürücü 1", D(47), D(99)), ("Sürücü 2", D(40), None)], {"kayit": 16, "ihlal": 6, "ihlalli_surucu": 1}, []))
            self.assertEqual(load_workbook(cikti).sheetnames, ["İhlaller", "Sürücü Özeti", "Haftalık Toplamlar", "Günlük Kayıtlar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "t.csv"
            y.write_text("Sürücü;Tarih;Sürüş Süresi;Kesintisiz Sürüş;Dinlenme\nA;05.10.2026;9,5;4:30;uzun\nA;05.10.2026;1;;\nA;dün;9;;\nA;06.10.2026;çok;;\nA;07.10.2026;3;4;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual([(k["tarih"].day, k["surus"], k["kesintisiz"], k["dinlenme"]) for k in l], [(5, D("9.5"), D("4.5"), None), (7, D(3), D(4), None)])
            self.assertEqual(sorted(u["tur"] for u in uy), ["Kayıt okunamadı", "Kayıt okunamadı", "Kesintisiz sürüş günlük sürüşten uzun", "Mükerrer gün", "Süre okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--kayitlar", str(y), "--cikti", c]
            self.assertEqual([u["tur"] for u in main.calistir(y, Path(c))["uyarilar"] if u["tur"] == "Eksik veri"], ["Eksik veri"])
            self.assertEqual(main.main(arg + ["--gunluk-surus", "8", "--uzatilmis-surus", "9:30", "--uzatma-hakki", "1", "--haftalik-surus", "48"]), 0)
            self.assertEqual(main.main(arg + ["--uzatilmis-surus", "8"]), 2)
            self.assertEqual(main.main(arg + ["--uzatma-hakki", "1,5"]), 2)
            self.assertEqual(main.main(["--kayitlar", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Sürücü;Tarih;Sürüş Süresi\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
