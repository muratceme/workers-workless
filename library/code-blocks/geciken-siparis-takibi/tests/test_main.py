import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
RAPOR = date(2026, 10, 12)          # pazartesi


def sp(no="S1", siparis=-3, termin=None, kargo=None, kanal="K", depo="D", durum=""):
    t = lambda g: None if g is None else RAPOR + timedelta(days=g)   # noqa: E731
    return {"no": no, "kanal": kanal, "depo": depo, "siparis": t(siparis), "termin": t(termin), "kargo": t(kargo), "durum": durum}


def tek(**kw):
    ayar = {k: kw.pop(k) for k in ("sure", "sure_gun", "is_gunu") if k in kw}
    return main.degerlendir([sp(**kw)], RAPOR, **ayar)["siparisler"][0]


class DurumTesti(unittest.TestCase):
    def test_bekleyen_siparisler(self):
        self.assertEqual([(x["sonuc"], x["kalan"], x["gecikme"]) for x in (tek(siparis=g) for g in (-5, -2, -1))], [("Gecikmiş", -3, 3), ("Bugün son gün", 0, None), ("Süresi var", 1, None)])
        x = tek(siparis=-6, termin=-4)
        self.assertEqual((x["sonuc"], x["gecikme"], x["termin_kaynak"]), ("Gecikmiş", 4, "Dosya"))
        self.assertEqual(tek(siparis=-2, sure={"k": 1})["sonuc"], "Gecikmiş")
        self.assertEqual(tek(siparis=-2, sure={"baska": 1})["sonuc"], "Bugün son gün")
        self.assertEqual(tek(siparis=-2, sure_gun=0)["gecikme"], 2)

    def test_kargolanan_ve_is_gunu(self):
        self.assertEqual([(x["sonuc"], x["gecikme"]) for x in (tek(siparis=-5, kargo=k) for k in (-4, -3, -1))], [("Zamanında kargolandı", None), ("Zamanında kargolandı", None), ("Geç kargolandı", 2)])
        cuma = date(2026, 10, 9)
        self.assertEqual([main.gun_ekle(cuma, n, True) for n in (0, 1, 2, 6)], [cuma, date(2026, 10, 12), date(2026, 10, 13), date(2026, 10, 19)])
        self.assertEqual(main.gun_ekle(date(2026, 10, 10), 1, True), date(2026, 10, 12))                # cumartesi siparişi + 1 iş günü = pazartesi
        self.assertEqual(main.gun_ekle(cuma, 2, False), date(2026, 10, 11))
        self.assertEqual((tek(siparis=-3, sure_gun=1)["sonuc"], tek(siparis=-3, sure_gun=1, is_gunu=True)["sonuc"]), ("Gecikmiş", "Bugün son gün"))

    def test_ozet_siralama_ve_iptal(self):
        l = [sp("A", -5), sp("B", -9, depo="E"), sp("C", -2), sp("D", -5, kargo=-4), sp("E", -5, kargo=-1, kanal="L"), sp("F", -5, kargo=-3, kanal="L"), sp("G", -9, durum="İptal edildi")]
        r = main.degerlendir(l, RAPOR)
        self.assertEqual([s["no"] for s in r["siparisler"]], ["B", "A", "C", "E", "D", "F"])
        self.assertEqual((r["say"], r["iptal"], r["ortalama"]), ({"Gecikmiş": 2, "Bugün son gün": 1, "Süresi var": 0, "Geç kargolandı": 1, "Zamanında kargolandı": 2}, 1, D("5.00")))
        self.assertEqual([(x["kanal"], x["toplam"], x["Gecikmiş"], x["zamaninda_oran"], x["ortalama"], x["en_uzun"]) for x in r["kanal"]], [("K", 4, 2, D("100.00"), D("5.00"), 7), ("L", 2, 0, D("50.00"), None, None)])
        self.assertEqual([(x["depo"], x["Gecikmiş"]) for x in r["depo"]], [("E", 1), ("D", 1)])
        with self.assertRaises(ValueError):
            main.degerlendir([sp(durum="iptal")], RAPOR)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "g.xlsx"
            r = main.calistir(ORNEK / "siparisler.csv", cikti, RAPOR, ORNEK / "kanal_sureleri.csv")
            self.assertEqual([(s["no"], s["sonuc"], s["gecikme"]) for s in r["siparisler"]],
                             [("S-1007", "Gecikmiş", 4), ("S-1001", "Gecikmiş", 3), ("S-1006", "Gecikmiş", 1), ("S-1002", "Bugün son gün", None), ("S-1010", "Bugün son gün", None),
                              ("S-1003", "Süresi var", None), ("S-1005", "Geç kargolandı", 2), ("S-1004", "Zamanında kargolandı", None), ("S-1008", "Zamanında kargolandı", None)])
            self.assertEqual({x["kanal"]: (x["toplam"], x["Gecikmiş"], x["zamaninda_oran"]) for x in r["kanal"]},
                             {"Web Sitesi": (3, 1, None), "Pazaryeri A": (4, 1, D("50.00")), "Pazaryeri B": (2, 1, D("100.00"))})
            self.assertEqual(({x["depo"]: (x["toplam"], x["Gecikmiş"]) for x in r["depo"]}, r["ortalama"], r["iptal"], r["uyarilar"]), ({"İstanbul": (6, 1), "Ankara": (3, 2)}, D("2.67"), 1, []))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Bekleyen Siparişler", "Kanal Özeti", "Depo Özeti", "Kargolananlar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, sr = Path(tmp) / "s.csv", Path(tmp) / "k.csv"
            y.write_text("Sipariş No;Kanal;Sipariş Tarihi;Termin;Kargo Tarihi\nA;K;10.10.2026;yarın;sonra\nA;K;10.10.2026;09.10.2026;\nB;L;dün;;\nC;M;2026-10-11 14:30:00;;\n", encoding="utf-8")
            sr.write_text("Kanal;Süre\nK;1\nL;hızlı\n", encoding="utf-8")
            l, sure, uy = main.oku(y, sr)
            self.assertEqual(([(s["no"], s["siparis"], s["termin"], s["kargo"], s["depo"]) for s in l], sure),
                             ([("A", date(2026, 10, 10), None, None, "Belirtilmemiş"), ("A", date(2026, 10, 10), date(2026, 10, 9), None, "Belirtilmemiş"), ("C", date(2026, 10, 11), None, None, "Belirtilmemiş")],
                              {"k": 1}))
            self.assertEqual(sorted(u["tur"] for u in uy), ["Kanal süresi okunamadı", "Kargoya veriliş tarihi okunamadı", "Mükerrer sipariş no", "Sipariş tarihi okunamadı", "Termin okunamadı",
                                                            "Termin sipariş tarihinden önce"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--siparisler", str(y), "--cikti", c, "--tarih", "12.10.2026"]
            r = main.calistir(y, Path(c), RAPOR, sr)
            self.assertEqual([u["kim"] for u in r["uyarilar"] if u["tur"] == "Kanalın süresi tanımlı değil"], ["M"])
            self.assertEqual(main.main(arg + ["--kanal-sure", str(sr), "--sure-gun", "3", "--is-gunu"]), 0)
            self.assertEqual(main.main(arg + ["--sure-gun", "1,5"]), 2)
            self.assertEqual(main.main(arg[:-1] + ["yarın"]), 2)
            self.assertEqual(main.main(["--siparisler", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Sipariş No;Sipariş Tarihi\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
