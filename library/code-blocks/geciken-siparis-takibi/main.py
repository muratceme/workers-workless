"""
Geciken Sipariş Takibi — Gerisi Bende kod bloğu
E-ticaret Operasyon ve Kargo › E-ticaret Operasyon Uzmanı

Kargoya verilmesi gereken süreyi aşan siparişleri kanal ve depo bazında listeler.
  Termin          = dosyadaki "Termin" tarihi; boşsa sipariş tarihi + kanalın kargoya verme süresi (kanal süre dosyası, yoksa --sure-gun, varsayılan 2 gün)
                    --is-gunu verilirse süre iş günü sayılır (cumartesi ve pazar atlanır; resmi tatiller atlanmaz)
  Gecikmiş        = kargoya verilmemiş ve rapor tarihi > termin;        gecikme = rapor tarihi − termin (takvim günü)
  Bugün son gün   = kargoya verilmemiş ve rapor tarihi = termin
  Süresi var      = kargoya verilmemiş ve rapor tarihi < termin
  Geç kargolandı  = kargoya veriliş tarihi > termin;                    gecikme = kargo tarihi − termin
  Zamanında kargolandı = kargoya veriliş tarihi ≤ termin
  Zamanında kargolama % = zamanında kargolanan ÷ kargolanan × 100
Kargoya verme süreleri kanalın kendi kuralıdır ve değişir; kod hiçbir kanalın süresini bilmez, siz verirsiniz. İptal edilen siparişler değerlendirilmez.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"no": ("siparis no", "siparis numarasi", "siparis"), "kanal": ("kanal", "satis kanali", "pazaryeri"), "depo": ("depo", "cikis deposu", "magaza"), "siparis": ("siparis tarihi", "tarih"),
         "termin": ("termin", "son kargo tarihi", "kargoya son teslim", "kargoya verilmesi gereken tarih"), "kargo": ("kargoya verilis tarihi", "kargo tarihi", "kargoya verilis", "sevk tarihi"),
         "durum": ("durum", "siparis durumu")}
SURE = {"kanal": SUTUN["kanal"], "sure": ("sure", "sure gun", "kargoya verme suresi", "gun")}
SIRA = ("Gecikmiş", "Bugün son gün", "Süresi var", "Geç kargolandı", "Zamanında kargolandı")


def gun_ekle(d: date, n: int, is_gunu: bool) -> date:
    if not is_gunu:
        return d + timedelta(days=n)
    while n > 0:
        d += timedelta(days=1)
        n -= d.weekday() < 5
    return d


def oku(yol: Path, sure_yol: Path | None = None) -> tuple[list[dict], dict, list[dict]]:
    uy, l, sure, gorulen = [], [], {}, set()
    for r in kayitlar(yol, SUTUN, ("no", "siparis")):
        no = metin(r["no"])
        if not no:
            continue
        kim = f"{no} (satır {r['_satir']})"
        st = tarih(r["siparis"])
        if st is None:
            uy.append(uyari("Yüksek", "Sipariş tarihi okunamadı", kim, f"'{metin(r['siparis'])}'. Sipariş değerlendirilmedi."))
            continue
        if katla(no) in gorulen:
            uy.append(uyari("Orta", "Mükerrer sipariş no", kim, "Aynı sipariş numarası birden fazla satırda; hepsi ayrı değerlendirildi (çok paketli sipariş olabilir)."))
        gorulen.add(katla(no))
        x = {"no": no, "kanal": metin(r["kanal"]) or "Belirtilmemiş", "depo": metin(r["depo"]) or "Belirtilmemiş", "siparis": st, "durum": metin(r["durum"])}
        for a, ad in (("termin", "Termin"), ("kargo", "Kargoya veriliş tarihi")):
            x[a] = tarih(r[a])
            if x[a] is None and metin(r[a]):
                uy.append(uyari("Yüksek", f"{ad} okunamadı", kim, f"'{metin(r[a])}'. Boş sayıldı" + ("; termin kanal süresinden hesaplandı." if a == "termin" else "; sipariş kargoya verilmemiş kabul edildi.")))
            elif x[a] and x[a] < st:
                uy.append(uyari("Orta", f"{ad} sipariş tarihinden önce", kim, f"{ad} {x[a]:%d.%m.%Y}, sipariş tarihi {st:%d.%m.%Y}. Tarihleri kontrol edin."))
        l.append(x)
    if sure_yol:
        for r in kayitlar(sure_yol, SURE, ("kanal", "sure")):
            k, s = metin(r["kanal"]), para(r["sure"])
            if not k:
                continue
            if s is None or s < 0 or s != int(s):
                uy.append(uyari("Yüksek", "Kanal süresi okunamadı", k, f"{sure_yol.name} satır {r['_satir']}: '{metin(r['sure'])}'. Bu kanal için varsayılan süre kullanıldı."))
                continue
            sure[katla(k)] = int(s)
    return l, sure, uy


def degerlendir(siparisler: list[dict], rapor: date, sure: dict | None = None, sure_gun: int = 2, is_gunu: bool = False) -> dict:
    sure = sure or {}
    iptal = [s for s in siparisler if "iptal" in katla(s["durum"])]
    l = [s for s in siparisler if "iptal" not in katla(s["durum"])]
    if not l:
        raise ValueError("Değerlendirilecek sipariş yok.")
    for s in l:
        s["termin_kaynak"] = "Dosya" if s["termin"] else "Hesaplanan"
        s["termin"] = s["termin"] or gun_ekle(s["siparis"], sure.get(katla(s["kanal"]), sure_gun), is_gunu)
        s["gecikme"] = s["kalan"] = None
        if s["kargo"]:
            s["sonuc"] = "Zamanında kargolandı" if s["kargo"] <= s["termin"] else "Geç kargolandı"
            s["gecikme"] = (s["kargo"] - s["termin"]).days if s["kargo"] > s["termin"] else None
        else:
            s["kalan"] = (s["termin"] - rapor).days
            s["sonuc"] = "Gecikmiş" if s["kalan"] < 0 else "Bugün son gün" if s["kalan"] == 0 else "Süresi var"
            s["gecikme"] = -s["kalan"] if s["kalan"] < 0 else None

    def ozet(alan: str) -> list[dict]:
        g = {}
        for s in l:
            x = g.setdefault(katla(s[alan]), {alan: s[alan], "toplam": 0, "gunler": []} | {t: 0 for t in SIRA})
            x["toplam"] += 1
            x[s["sonuc"]] += 1
            if s["sonuc"] == "Gecikmiş":
                x["gunler"].append(s["gecikme"])
        for x in g.values():
            kargolanan = x["Zamanında kargolandı"] + x["Geç kargolandı"]
            x["zamaninda_oran"] = yuvarla(Decimal(x["Zamanında kargolandı"]) / kargolanan * 100) if kargolanan else None
            x["ortalama"] = yuvarla(Decimal(sum(x["gunler"])) / len(x["gunler"])) if x["gunler"] else None
            x["en_uzun"] = max(x["gunler"], default=None)
        return sorted(g.values(), key=lambda x: (-x["Gecikmiş"], -(x["en_uzun"] or 0), x[alan]))

    gunler = [s["gecikme"] for s in l if s["sonuc"] == "Gecikmiş"]
    return {"siparisler": sorted(l, key=lambda s: (SIRA.index(s["sonuc"]), -(s["gecikme"] or 0), s["termin"], s["no"])), "kanal": ozet("kanal"), "depo": ozet("depo"), "rapor": rapor, "iptal": len(iptal),
            "say": {t: sum(s["sonuc"] == t for s in l) for t in SIRA}, "ortalama": yuvarla(Decimal(sum(gunler)) / len(gunler)) if gunler else None, "ayar": (sure_gun, is_gunu),
            "varsayilan_sure": sorted({s["kanal"] for s in l if s["termin_kaynak"] == "Hesaplanan" and katla(s["kanal"]) not in sure}) if sure else []}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    renk = {"Gecikmiş": KIRMIZI, "Bugün son gün": SARI, "Geç kargolandı": MAVI}
    sg, ig = r["ayar"]
    ws = wb.active
    ws.title = "Bekleyen Siparişler"
    _baslik(ws, ["Durum", "Sipariş No", "Kanal", "Depo", "Sipariş Tarihi", "Termin", "Termin Kaynağı", "Kalan Gün", "Gecikme (gün)", "Sipariş Durumu", "Not (siz doldurun)"], (16, 18, 20, 18, 12, 12, 12, 9, 10, 18, 30))
    for s in r["siparisler"]:
        if not s["kargo"]:
            _satir(ws, [s["sonuc"], s["no"], s["kanal"], s["depo"], s["siparis"], s["termin"], s["termin_kaynak"], s["kalan"], s["gecikme"], s["durum"], ""], tarih_s=(5, 6), dolgu={1: renk.get(s["sonuc"])})
    ws.append([])
    ws.append([f"Rapor tarihi {r['rapor']:%d.%m.%Y}. Termini boş siparişlerde süre kanal dosyasından, yoksa {sg} gün ({'iş günü' if ig else 'takvim günü'}). Gecikme takvim günüdür. "
               f"İptal edilen {r['iptal']} sipariş değerlendirilmedi."])
    for ad, alan, bas in (("Kanal Özeti", "kanal", "Kanal"), ("Depo Özeti", "depo", "Depo")):
        ws = wb.create_sheet(ad)
        _baslik(ws, [bas, "Sipariş", *SIRA, "Zamanında Kargolama %", "Ortalama Gecikme (gün)", "En Uzun Gecikme (gün)"], (22, 9, 10, 10, 10, 10, 12, 12, 12, 12))
        for x in r[alan]:
            _satir(ws, [x[alan], x["toplam"], *[x[t] or None for t in SIRA], x["zamaninda_oran"], x["ortalama"], x["en_uzun"]], para_s=(8, 9), dolgu={3: KIRMIZI if x["Gecikmiş"] else None})
    ws = wb.create_sheet("Kargolananlar")
    _baslik(ws, ["Sonuç", "Sipariş No", "Kanal", "Depo", "Sipariş Tarihi", "Termin", "Kargoya Veriliş", "Gecikme (gün)"], (20, 18, 20, 18, 12, 12, 12, 10))
    for s in r["siparisler"]:
        if s["kargo"]:
            _satir(ws, [s["sonuc"], s["no"], s["kanal"], s["depo"], s["siparis"], s["termin"], s["kargo"], s["gecikme"]], tarih_s=(5, 6, 7), dolgu={1: renk.get(s["sonuc"])})
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, rapor: date, sure_yol: Path | None = None, sure_gun: int = 2, is_gunu: bool = False) -> dict:
    siparisler, sure, uy = oku(yol, sure_yol)
    r = degerlendir(siparisler, rapor, sure, sure_gun, is_gunu)
    for k in r["varsayilan_sure"]:
        uy.append(uyari("Orta", "Kanalın süresi tanımlı değil", k, f"Kanal süre dosyasında yok; varsayılan {sure_gun} gün kullanıldı."))
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Kargoya verilme süresini aşan siparişleri kanal ve depo bazında listeler.")
    p.add_argument("--siparisler", type=Path, default=ORNEK / "siparisler.csv", help="Sipariş No, Kanal, Depo, Sipariş Tarihi, Termin, Kargoya Veriliş Tarihi, Durum")
    p.add_argument("--kanal-sure", type=Path, help="İsteğe bağlı: Kanal, Süre (gün)")
    p.add_argument("--sure-gun", default="2", help="Süresi tanımlı olmayan kanallar için kargoya verme süresi (varsayılan 2 gün)")
    p.add_argument("--is-gunu", action="store_true", help="Süreyi iş günü say (cumartesi ve pazar atlanır; resmi tatiller atlanmaz)")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 12.10.2026)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "geciken_siparisler.xlsx")
    a = p.parse_args(argv)
    ornek = a.siparisler == ORNEK / "siparisler.csv"
    sure_yol = a.kanal_sure or (ORNEK / "kanal_sureleri.csv" if ornek else None)
    rapor, sg = tarih(a.tarih) if a.tarih else (date(2026, 10, 12) if ornek else date.today()), para(a.sure_gun)
    if rapor is None or sg is None or sg < 0 or sg != int(sg):
        print("[X] --tarih GG.AA.YYYY biçiminde; --sure-gun sıfır veya pozitif tam sayı olmalı")
        return 2
    for y in (a.siparisler, sure_yol):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.siparisler, a.cikti, rapor, sure_yol, int(sg), a.is_gunu)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    s = r["say"]
    print(f"[OK] {len(r['siparisler'])} sipariş · {s['Gecikmiş']} gecikmiş, {s['Bugün son gün']} bugün son gün, {s['Süresi var']} süresi var · {s['Geç kargolandı']} geç, {s['Zamanında kargolandı']} zamanında kargolandı")
    for x in r["siparisler"]:
        if x["sonuc"] == "Gecikmiş":
            print(f"[!] {x['no']} · {x['kanal']} · {x['depo']}: termin {x['termin']:%d.%m.%Y}, {x['gecikme']} gün gecikmiş")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
