"""
Yakıt Tüketim Analizi — Workers / Workless kod bloğu
Filo Yönetimi › Filo Yöneticisi

Yakıt kartı ve kilometre verisinden araç ve sürücü bazında 100 km tüketimini hesaplayıp sapmaları işaretler.
  Dilim             = aynı aracın art arda iki yakıt alımı arası: mesafe = km sayacı farkı, yakıt = ikinci alımın litresi   (her alımda depo doldurulduğu varsayılır)
  Araç tüketimi     = (ilk alım hariç litre toplamı) ÷ (son km − ilk km) × 100                                              (L / 100 km)
  Km başına maliyet = (ilk alım hariç tutar toplamı) ÷ toplam mesafe
  Hedefe göre sapma = (araç tüketimi − hedef) ÷ hedef × 100;  sapma > --tolerans (%10) ise "Hedefin üzerinde"                (hedef araç dosyasından)
  Şüpheli dilim     = dilim tüketimi, aracın ortalamasının --dilim-tolerans (%30) üzerinde
  Diğer kontroller  = km sayacı geri gitmiş veya değişmemiş · alınan litre depo kapasitesinden fazla
  Sürücü tüketimi   = sürücünün yakıt aldığı dilimlerin litre toplamı ÷ mesafe toplamı × 100                                 (dilim, alımı yapan sürücüye yazılır)
Depo her alımda doldurulmuyorsa dilim tüketimleri yanıltır; araç toplamı uzun dönemde yine de anlamlıdır. Sapma suistimal kanıtı değildir: yük, güzergah, mevsim ve rölanti etkiler.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"plaka": ("plaka", "arac", "arac plakasi"), "tarih": ("tarih", "islem tarihi", "alim tarihi"), "km": ("km", "km sayaci", "kilometre", "sayac"), "litre": ("litre", "miktar", "lt"),
         "tutar": ("tutar", "toplam tutar", "tl"), "surucu": ("surucu", "sofor", "kart sahibi")}
ARAC = {"plaka": SUTUN["plaka"], "hedef": ("hedef tuketim", "hedef tuketim l 100 km", "hedef", "hedef l 100 km", "beklenen tuketim"), "depo": ("depo kapasitesi", "depo kapasitesi l", "depo", "depo litre"), "tip": ("arac tipi", "tip", "model")}
Z = Decimal(0)


def plaka_anahtar(x) -> str:
    return katla(x).replace(" ", "")


def oku(yol: Path, arac_yol: Path | None = None) -> tuple[list[dict], dict, list[dict]]:
    uy, l, arac = [], [], {}
    for r in kayitlar(yol, SUTUN, ("plaka", "tarih", "km", "litre")):
        p = metin(r["plaka"])
        if not p:
            continue
        kim = f"{p} (satır {r['_satir']})"
        t, km, litre, tutar = tarih(r["tarih"]), para(r["km"]), para(r["litre"]), para(r["tutar"])
        if t is None or km is None or litre is None or km < 0 or litre <= 0:
            uy.append(uyari("Yüksek", "Alım okunamadı", kim, f"Tarih '{metin(r['tarih'])}', km '{metin(r['km'])}', litre '{metin(r['litre'])}'. Satır alınmadı; komşu dilimlerin tüketimi yanlış çıkabilir."))
            continue
        l.append({"plaka": p, "tarih": t, "km": km, "litre": litre, "tutar": tutar if tutar and tutar > 0 else Z, "surucu": metin(r["surucu"]) or "Belirtilmemiş", "satir": r["_satir"]})
    if arac_yol:
        for r in kayitlar(arac_yol, ARAC, ("plaka",)):
            p = metin(r["plaka"])
            if not p:
                continue
            h, d = para(r["hedef"]), para(r["depo"])
            if (h is None and metin(r["hedef"])) or (d is None and metin(r["depo"])) or (h is not None and h <= 0) or (d is not None and d <= 0):
                uy.append(uyari("Orta", "Araç bilgisi okunamadı", p, f"Hedef '{metin(r['hedef'])}', depo '{metin(r['depo'])}'. Okunamayan alan boş sayıldı."))
            arac[plaka_anahtar(p)] = {"hedef": h if h and h > 0 else None, "depo": d if d and d > 0 else None, "tip": metin(r["tip"])}
    return l, arac, uy


def analiz(alimlar: list[dict], arac: dict | None = None, tolerans: Decimal = Decimal(10), dilim_tolerans: Decimal = Decimal(30)) -> dict:
    if not alimlar:
        raise ValueError("Analiz edilecek yakıt alımı yok.")
    arac, grup, uy, dilimler, araclar = arac or {}, defaultdict(list), [], [], []
    for a in alimlar:
        grup[plaka_anahtar(a["plaka"])].append(a)
    for k, liste in grup.items():
        liste.sort(key=lambda a: (a["tarih"], a["km"]))
        bilgi, p = arac.get(k, {"hedef": None, "depo": None, "tip": ""}), liste[0]["plaka"]
        kendi, onceki = [], liste[0]
        for a in liste:
            if bilgi["depo"] and a["litre"] > bilgi["depo"]:
                uy.append(uyari("Yüksek", "Depo kapasitesini aşan alım", p, f"{a['tarih']:%d.%m.%Y}: {a['litre']:f} L alınmış, depo kapasitesi {bilgi['depo']:f} L."))
        for a in liste[1:]:
            mesafe = a["km"] - onceki["km"]
            if mesafe <= 0:
                uy.append(uyari("Yüksek", "Km sayacı ilerlememiş", p, f"{a['tarih']:%d.%m.%Y}: km {a['km']:f}, önceki alımda {onceki['km']:f}. Bu alım dilim hesabına girmedi; km girişini kontrol edin."))
                continue
            kendi.append({"plaka": p, "tarih": a["tarih"], "onceki": onceki["tarih"], "surucu": a["surucu"], "mesafe": mesafe, "litre": a["litre"], "tutar": a["tutar"], "tuketim": yuvarla(a["litre"] / mesafe * 100)})
            onceki = a
        if not kendi:
            uy.append(uyari("Orta", "Tüketim hesaplanamadı", p, "Hesap için aynı araca ait, km sayacı ilerlemiş en az iki alım gerekir."))
            continue
        mesafe, litre, tutar = sum((d["mesafe"] for d in kendi), Z), sum((d["litre"] for d in kendi), Z), sum((d["tutar"] for d in kendi), Z)
        ort = litre / mesafe * 100
        for d in kendi:
            d["suphe"] = len(kendi) > 1 and d["tuketim"] > ort * (1 + dilim_tolerans / 100)
        sapma = yuvarla((ort - bilgi["hedef"]) / bilgi["hedef"] * 100) if bilgi["hedef"] else None
        araclar.append({"plaka": p, "tip": bilgi["tip"], "alim": len(liste), "mesafe": mesafe, "litre": litre, "tutar": tutar, "tuketim": yuvarla(ort), "km_maliyet": yuvarla(tutar / mesafe) if tutar else None,
                        "hedef": bilgi["hedef"], "sapma": sapma, "durum": "Hedef yok" if sapma is None else "Hedefin üzerinde" if sapma > tolerans else "Hedefte", "suphe": sum(d["suphe"] for d in kendi)})
        dilimler += kendi
    if not araclar:
        raise ValueError("Hiçbir araç için tüketim hesaplanamadı: aynı araca ait en az iki alım gerekir.")
    sur = {}
    for d in dilimler:
        s = sur.setdefault(katla(d["surucu"]), {"surucu": d["surucu"], "dilim": 0, "mesafe": Z, "litre": Z, "suphe": 0, "araclar": set()})
        s["dilim"] += 1
        s["mesafe"] += d["mesafe"]
        s["litre"] += d["litre"]
        s["suphe"] += d["suphe"]
        s["araclar"].add(d["plaka"])
    for s in sur.values():
        s["tuketim"] = yuvarla(s["litre"] / s["mesafe"] * 100)
    m, lt = sum((a["mesafe"] for a in araclar), Z), sum((a["litre"] for a in araclar), Z)
    return {"araclar": sorted(araclar, key=lambda a: (-(a["sapma"] if a["sapma"] is not None else Decimal(-10**6)), a["plaka"])), "dilimler": sorted(dilimler, key=lambda d: (not d["suphe"], d["plaka"], d["tarih"])),
            "suruculer": sorted(sur.values(), key=lambda s: -s["tuketim"]), "filo": {"mesafe": m, "litre": lt, "tutar": sum((a["tutar"] for a in araclar), Z), "tuketim": yuvarla(lt / m * 100)}, "uyarilar": uy,
            "ayar": (tolerans, dilim_tolerans), "say": {"ustunde": sum(a["durum"] == "Hedefin üzerinde" for a in araclar), "suphe": sum(d["suphe"] for d in dilimler)}}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    tol, dtol = r["ayar"]
    f = r["filo"]
    ws = wb.active
    ws.title = "Araç Bazında"
    _baslik(ws, ["Plaka", "Araç Tipi", "Alım Sayısı", "Mesafe (km)", "Yakıt (L)", "Tüketim (L/100 km)", "Hedef", "Sapma %", "Durum", "Yakıt Tutarı", "Km Başına Maliyet", "Şüpheli Dilim", "Not (siz doldurun)"],
            (14, 18, 8, 12, 11, 11, 8, 9, 16, 14, 11, 9, 30))
    for a in r["araclar"]:
        _satir(ws, [a["plaka"], a["tip"], a["alim"], a["mesafe"], a["litre"], a["tuketim"], a["hedef"], a["sapma"], a["durum"], a["tutar"] or None, a["km_maliyet"], a["suphe"] or None, ""],
               para_s=(5, 6, 7, 8, 10, 11), dolgu={9: KIRMIZI if a["durum"] == "Hedefin üzerinde" else YESIL if a["durum"] == "Hedefte" else None, 12: SARI if a["suphe"] else None})
    _satir(ws, ["Filo", "", None, f["mesafe"], f["litre"], f["tuketim"], None, None, "", f["tutar"] or None], para_s=(5, 6, 10))
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.append([])
    ws.append([f"Tüketim ilk alım hariç litre ÷ mesafe × 100. Hedefe göre sapma > %{tol.normalize():f} ise 'Hedefin üzerinde'; dilim tüketimi araç ortalamasının %{dtol.normalize():f} üzerindeyse 'şüpheli'. "
               "Her alımda depo doldurulmuyorsa dilimler yanıltır. Sapma suistimal kanıtı değildir."])
    ws = wb.create_sheet("Sürücü Bazında")
    _baslik(ws, ["Sürücü", "Dilim Sayısı", "Araç Sayısı", "Mesafe (km)", "Yakıt (L)", "Tüketim (L/100 km)", "Şüpheli Dilim"], (26, 9, 9, 12, 11, 11, 9))
    for s in r["suruculer"]:
        _satir(ws, [s["surucu"], s["dilim"], len(s["araclar"]), s["mesafe"], s["litre"], s["tuketim"], s["suphe"] or None], para_s=(5, 6), dolgu={7: SARI if s["suphe"] else None})
    ws = wb.create_sheet("Dilimler")
    _baslik(ws, ["Şüpheli", "Plaka", "Önceki Alım", "Alım Tarihi", "Sürücü", "Mesafe (km)", "Yakıt (L)", "Tüketim (L/100 km)", "Tutar"], (9, 14, 12, 12, 24, 12, 11, 11, 14))
    for d in r["dilimler"]:
        _satir(ws, ["Evet" if d["suphe"] else "", d["plaka"], d["onceki"], d["tarih"], d["surucu"], d["mesafe"], d["litre"], d["tuketim"], d["tutar"] or None], para_s=(7, 8, 9), tarih_s=(3, 4),
               dolgu={1: SARI if d["suphe"] else None})
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, arac_yol: Path | None = None, **ayar) -> dict:
    alimlar, arac, uy = oku(yol, arac_yol)
    r = analiz(alimlar, arac, **ayar)
    uy += r["uyarilar"]
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Yakıt alımlarından araç ve sürücü bazında 100 km tüketimini hesaplar.")
    p.add_argument("--alimlar", type=Path, default=ORNEK / "yakit_alimlari.csv", help="Plaka, Tarih, Km (sayaç), Litre, Tutar, Sürücü")
    p.add_argument("--araclar", type=Path, help="İsteğe bağlı: Plaka, Hedef Tüketim (L/100 km), Depo Kapasitesi (L), Araç Tipi")
    p.add_argument("--tolerans", default="10", help="Hedefin bu %%'den fazla üzerindeki araç işaretlenir (varsayılan 10)")
    p.add_argument("--dilim-tolerans", default="30", help="Araç ortalamasının bu %%'den fazla üzerindeki dilim şüpheli sayılır (varsayılan 30)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "yakit_tuketim.xlsx")
    a = p.parse_args(argv)
    arac = a.araclar or (ORNEK / "araclar.csv" if a.alimlar == ORNEK / "yakit_alimlari.csv" else None)
    t, dt = para(a.tolerans), para(a.dilim_tolerans)
    if t is None or dt is None or t < 0 or dt < 0:
        print("[X] --tolerans ve --dilim-tolerans sıfır veya pozitif sayı olmalı")
        return 2
    for y in (a.alimlar, arac):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.alimlar, a.cikti, arac, tolerans=t, dilim_tolerans=dt)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    f, s = r["filo"], r["say"]
    print(f"[OK] {len(r['araclar'])} araç · {f['mesafe']:,.0f} km · {f['litre']:,.2f} L · filo tüketimi {f['tuketim']} L/100 km · {s['ustunde']} araç hedefin üzerinde, {s['suphe']} şüpheli dilim")
    for x in r["araclar"]:
        if x["durum"] == "Hedefin üzerinde":
            print(f"[!] {x['plaka']}: {x['tuketim']} L/100 km, hedef {x['hedef']:f} (sapma %{x['sapma']})")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
