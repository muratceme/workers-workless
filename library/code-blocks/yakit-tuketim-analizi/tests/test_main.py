import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def al(gun, km, litre, surucu="S", plaka="34 A 1", tutar=None):
    return {"plaka": plaka, "tarih": date(2026, 10, gun), "km": D(km), "litre": D(litre), "tutar": D(litre) * 40 if tutar is None else D(tutar), "surucu": surucu, "satir": 0}


def bilgi(hedef=None, depo=None):
    return {"34a1": {"hedef": None if hedef is None else D(hedef), "depo": None if depo is None else D(depo), "tip": ""}}


class AnalizTesti(unittest.TestCase):
    def test_arac_tuketimi(self):
        l = [al(1, 100000, 400, "Ali"), al(4, 101000, 300, "Ali"), al(7, 102000, 310, "Ali"), al(10, 103000, 480, "Veli")]
        r = main.analiz(l, bilgi(30, 600))
        a = r["araclar"][0]
        self.assertEqual((a["mesafe"], a["litre"], a["tuketim"], a["km_maliyet"], a["sapma"], a["durum"], a["suphe"], a["alim"]), (D(3000), D(1090), D("36.33"), D("14.53"), D("21.11"), "Hedefin üzerinde", 1, 4))
        self.assertEqual([(d["tuketim"], d["suphe"], d["surucu"]) for d in r["dilimler"]], [(D("48.00"), True, "Veli"), (D("30.00"), False, "Ali"), (D("31.00"), False, "Ali")])      # ilk alım yalnız başlangıç
        self.assertEqual([(s["surucu"], s["mesafe"], s["litre"], s["tuketim"], s["suphe"]) for s in r["suruculer"]], [("Veli", D(1000), D(480), D("48.00"), 1), ("Ali", D(2000), D(610), D("30.50"), 0)])
        self.assertEqual([main.analiz([al(1, 0, 1), al(2, 1000, lt)], bilgi(30))["araclar"][0]["durum"] for lt in (330, 331)], ["Hedefte", "Hedefin üzerinde"])                       # tam %10 hedefte
        self.assertEqual((main.analiz([al(1, 0, 1), al(2, 1000, 400)])["araclar"][0]["durum"], main.analiz([al(1, 0, 1), al(2, 1000, 400)], bilgi(30), D(40))["araclar"][0]["durum"]), ("Hedef yok", "Hedefte"))
        self.assertEqual(main.analiz([al(2, 1000, 300), al(1, 0, 50)])["araclar"][0]["tuketim"], D("30.00"))                                                                         # tarihe göre sıralanır

    def test_kontroller(self):
        r = main.analiz([al(1, 1000, 100), al(2, 900, 100), al(3, 2000, 300), al(4, 2000, 50), al(5, 3000, 700)], bilgi(depo=600))
        self.assertEqual((sorted(u["tur"] for u in r["uyarilar"]), r["araclar"][0]["mesafe"], r["araclar"][0]["litre"]), (["Depo kapasitesini aşan alım", "Km sayacı ilerlememiş", "Km sayacı ilerlememiş"], D(2000), D(1000)))
        r = main.analiz([al(1, 0, 1), al(2, 1000, 300), al(1, 0, 1, plaka="B"), al(1, 0, 1, plaka="C"), al(2, 500, 200, plaka="C")])
        self.assertEqual(([a["plaka"] for a in r["araclar"]], r["filo"], [u["tur"] for u in r["uyarilar"]]),
                         (["34 A 1", "C"], {"mesafe": D(1500), "litre": D(500), "tutar": D(20000), "tuketim": D("33.33")}, ["Tüketim hesaplanamadı"]))
        self.assertEqual(main.analiz([al(1, 0, 1), al(2, 1000, 300, plaka="34a1")])["araclar"][0]["mesafe"], D(1000))                 # plaka yazım farkı birleştirilir
        for kotu in ([], [al(1, 0, 1)]):
            with self.assertRaises(ValueError):
                main.analiz(kotu)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "y.xlsx"
            r = main.calistir(ORNEK / "yakit_alimlari.csv", cikti, ORNEK / "araclar.csv")
            self.assertEqual([(a["plaka"], a["tuketim"], a["sapma"], a["durum"], a["km_maliyet"]) for a in r["araclar"]],
                             [("34 ORN 001", D("36.33"), D("21.11"), "Hedefin üzerinde", D("14.53")), ("34 ORN 002", D("31.00"), D("-3.13"), "Hedefte", D("12.40"))])
            self.assertEqual((r["filo"]["tuketim"], r["say"], r["uyarilar"]), (D("34.08"), {"ustunde": 1, "suphe": 1}, []))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Araç Bazında", "Sürücü Bazında", "Dilimler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, a = Path(tmp) / "y.csv", Path(tmp) / "a.csv"
            y.write_text("Plaka;Tarih;Km;Litre\n34 A 1;01.10.2026;100.000;400,50\n34 A 1;dün;1;1\n34 A 1;02.10.2026;101.000;0\n34A1;03.10.2026;101.500;450\n", encoding="utf-8")
            a.write_text("Plaka;Hedef Tüketim;Depo Kapasitesi\n34 A 1;30;çok\n", encoding="utf-8")
            l, ar, uy = main.oku(y, a)
            self.assertEqual(([(x["km"], x["litre"], x["tutar"], x["surucu"]) for x in l], ar, sorted(u["tur"] for u in uy)),
                             ([(D(100000), D("400.50"), D(0), "Belirtilmemiş"), (D(101500), D(450), D(0), "Belirtilmemiş")], {"34a1": {"hedef": D(30), "depo": None, "tip": ""}}, ["Alım okunamadı", "Alım okunamadı", "Araç bilgisi okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--alimlar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--araclar", str(a), "--tolerans", "5", "--dilim-tolerans", "20"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(arg + ["--tolerans", "-1"]), 2)
            self.assertEqual(main.main(arg + ["--araclar", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Plaka;Tarih;Km;Litre\nA;01.10.2026;1;1\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
