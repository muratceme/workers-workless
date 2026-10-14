import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
KADEME = [(D(90), D(1)), (D(100), D(2)), (D(120), D(3))]


def s(personel, gun, ciro, fis, adet):
    return {"tarih": date(2026, 10, gun) if gun else None, "personel": personel, "ciro": D(ciro), "fis": D(fis), "adet": D(adet)}


def veri():
    return [s("A", 1, 10000, 50, 120), s("a", 2, 12000, 60, 150), s("B", 1, 8000, 40, 80), s("C", None, 5000, 50, 100)], {"a": ("A", D(20000)), "b": ("B", D(10000)), "d": ("D", D(5000))}


class DegerlendirTesti(unittest.TestCase):
    def test_performans_ve_prim(self):
        r = main.degerlendir(*veri(), KADEME)
        self.assertEqual([(p["sira"], p["personel"], p["ciro"], p["pay"], p["sepet"], p["sepet_endeks"], p["upt"], p["upt_endeks"], p["gun"], p["gunluk"], p["gerceklesme"], p["prim_oran"], p["prim"]) for p in r["personel"]],
                         [(1, "A", D(22000), D("62.86"), D("200.00"), D("114.29"), D("2.45"), D("109.09"), 2, D("11000.00"), D("110.00"), D(2), D("440.00")),
                          (2, "B", D(8000), D("22.86"), D("200.00"), D("114.29"), D("2.00"), D("88.89"), 1, D("8000.00"), D("80.00"), D(0), D("0.00")),
                          (3, "C", D(5000), D("14.29"), D("100.00"), D("57.14"), D("2.00"), D("88.89"), None, None, None, None, None)])
        self.assertEqual(r["magaza"], {"ciro": D(35000), "fis": D(200), "adet": D(450), "sepet": D("175.00"), "upt": D("2.25"), "kisi": 3, "prim": D("440.00")})
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]), [("Bilgi", "Hedefi tanımlı değil", "C"), ("Orta", "Hedefi var, satışı yok", "D")])

    def test_kademe_ve_esik(self):
        p = lambda ciro, kademe=KADEME: main.degerlendir([s("A", 1, ciro, 1, 1)], {"a": ("A", D(1000))}, kademe)["personel"][0]   # noqa: E731
        self.assertEqual([(p(c)["prim_oran"], p(c)["prim"]) for c in (899, 900, 999, 1000, 1199, 1200, 2000)],
                         [(D(0), D("0.00")), (D(1), D("9.00")), (D(1), D("9.99")), (D(2), D("20.00")), (D(2), D("23.98")), (D(3), D("36.00")), (D(3), D("60.00"))])
        self.assertEqual((p(1000, None)["prim"], main.degerlendir(*veri())["magaza"]["prim"], main.degerlendir(veri()[0])["magaza"]["prim"]), (None, None, None))
        u = lambda esik: [x["tur"] for x in main.degerlendir([s("A", 1, 790, 1, 1)], {"a": ("A", D(1000))}, KADEME, D(esik))["uyarilar"]]   # noqa: E731
        self.assertEqual((u(79), u(80)), ([], ["Hedefin gerisinde"]))
        self.assertEqual([main.kademe_oku(x) for x in ("90:1;100:2", "95:0,5", "100:2;90:1", "90:1;90:2", "90", "x:1", "")], [[(D(90), D(1)), (D(100), D(2))], [(D(95), D("0.5"))], None, None, None, None, None])
        with self.assertRaises(ValueError):
            main.degerlendir([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "personel_satislari.csv", cikti, ORNEK / "hedefler.csv", KADEME)
            self.assertEqual([(p["personel"], p["ciro"], p["pay"], p["sepet"], p["sepet_endeks"], p["gun"], p["gerceklesme"], p["prim_oran"], p["prim"]) for p in r["personel"]],
                             [("Zeynep A.", D(132000), D("34.92"), D("400.00"), D("146.03"), 6, D("120.00"), D(3), D("3960.00")), ("Ayşe K.", D(108000), D("28.57"), D("300.00"), D("109.52"), 6, D("108.00"), D(2), D("2160.00")),
                              ("Mehmet T.", D(84000), D("22.22"), D("200.00"), D("73.02"), 6, D("93.33"), D(1), D("840.00")), ("Can B.", D(54000), D("14.29"), D("200.00"), D("73.02"), 6, D("67.50"), D(0), D("0.00"))])
            self.assertEqual(r["magaza"], {"ciro": D(378000), "fis": D(1380), "adet": D(3162), "sepet": D("273.91"), "upt": D("2.29"), "kisi": 4, "prim": D("6960.00")})
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]], [("Bilgi", "Hedefin gerisinde", "Can B.")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Personel", "Prim Tablosu", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, h = Path(tmp) / "s.csv", Path(tmp) / "h.csv"
            y.write_text("Gün;Satış Danışmanı;Satış;Fiş Sayısı;Adet\n01.10.2026;Ali;12.500,50;1.200;2.400\n;Veli;100;;\nx;Ali;5;;\n02.10.2026;Ali;y;;\n02.10.2026;Ali;5;-1;\n02.10.2026;;5;;\n", encoding="utf-8")
            h.write_text("Çalışan;Aylık Hedef\nAli;10.000\nali;5\nVeli;0\nCan;x\n", encoding="utf-8")
            l, uy = main.oku_satis(y)
            self.assertEqual(([(x["tarih"], x["personel"], x["ciro"], x["fis"], x["adet"]) for x in l], [u["tur"] for u in uy]),
                             ([(date(2026, 10, 1), "Ali", D("12500.50"), D(1200), D(2400)), (None, "Veli", D(100), D(0), D(0))], ["Satış satırı okunamadı"] * 3))
            hd, uy = main.oku_hedef(h)
            self.assertEqual((hd, [(u["tur"], u["kim"]) for u in uy]), ({"ali": ("Ali", D(10000))}, [("Hedef satırı okunamadı", "ali"), ("Hedef satırı okunamadı", "Veli"), ("Hedef satırı okunamadı", "Can")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--satislar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(load_workbook(c).sheetnames, ["Personel", "Uyarılar"])
            self.assertEqual(main.main(arg + ["--hedefler", str(h), "--kademeler", "95:0,5;100:1,5", "--esik", "85"]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--kademeler", "100:1;90:2"], ["--kademeler", "x"], ["--esik", "-1"])], [2] * 3)
            self.assertEqual((main.main(["--satislar", str(Path(tmp) / "yok.csv")]), main.main(arg + ["--hedefler", str(Path(tmp) / "yok.csv")])), (1, 1))
            for icerik in ("Personel;Fiş Adedi\nA;5\n", "Personel;Ciro\n;5\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
