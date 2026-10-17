"""
Personel Satış Performansı — Gerisi Bende kod bloğu
Mağaza Operasyonları › Mağaza Müdürü

Satış danışmanlarının ciro, fiş, adet ve sepet ortalamalarını karşılaştırır; hedef ve prim kademeleri verilirse prim tablosunu hazırlar.
  Sepet ortalaması       = ciro ÷ fiş adedi;   fiş başına adet = satılan adet ÷ fiş adedi
  Ciro payı %            = personelin cirosu ÷ mağaza cirosu × 100
  Mağazaya göre %        = personelin sepet ortalaması (ya da fiş başına adedi) ÷ mağazanınki × 100
  Hedef gerçekleşme %    = ciro ÷ hedef ciro × 100
  Prim oranı             = gerçekleşmenin ulaştığı en yüksek kademe   (--kademeler 90:1;100:2;120:3 → %90'da cironun %1'i, %100'de %2'si …)
  Prim tutarı            = ciro × prim oranı ÷ 100
Kademeler ve hedefler sizin prim politikanızdır; kod varsayılan oran kullanmaz. Vergi, SGK ve bordro hesabı yapılmaz.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

S_SUTUN = {"tarih": ("tarih", "gun"), "personel": ("personel", "satis danismani", "calisan", "ad soyad"), "ciro": ("ciro", "satis", "satis tutari", "net satis"), "fis": ("fis adedi", "fis sayisi", "islem adedi"),
           "adet": ("satilan adet", "adet", "urun adedi")}
H_SUTUN = {"personel": ("personel", "satis danismani", "calisan", "ad soyad"), "hedef": ("hedef ciro", "hedef", "aylik hedef")}
SIFIR = Decimal(0)


def oku_satis(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, S_SUTUN, ("personel", "ciro")):
        personel = metin(r["personel"])
        if not personel:
            continue
        d = {a: para(r[a]) if metin(r[a]) else SIFIR for a in ("ciro", "fis", "adet")}
        if any(v is None for v in d.values()) or d["fis"] < 0 or d["adet"] < 0 or (metin(r["tarih"]) and tarih(r["tarih"]) is None):
            uy.append(uyari("Yüksek", "Satış satırı okunamadı", personel, f"{yol.name} satır {r['_satir']}: tarih ya da sayılardan biri okunamadı. Satır alınmadı."))
            continue
        l.append({"tarih": tarih(r["tarih"]), "personel": personel, **d})
    return l, uy


def oku_hedef(yol: Path) -> tuple[dict, list[dict]]:
    uy, d = [], {}
    for r in kayitlar(yol, H_SUTUN, ("personel", "hedef")):
        personel = metin(r["personel"])
        if not personel:
            continue
        hedef = para(r["hedef"])
        if hedef is None or hedef <= 0 or katla(personel) in d:
            uy.append(uyari("Yüksek", "Hedef satırı okunamadı", personel, f"{yol.name} satır {r['_satir']}: " + ("aynı personel daha önce geçti" if katla(personel) in d else f"hedef '{metin(r['hedef'])}' okunamadı") + ". Satır alınmadı."))
            continue
        d[katla(personel)] = (personel, hedef)
    return d, uy


def kademe_oku(metni: str) -> list[tuple[Decimal, Decimal]] | None:
    """'90:1;100:2;120:3' → [(90, 1), (100, 2), (120, 3)]"""
    l = []
    for parca in metni.split(";"):
        if ":" not in parca:
            return None
        e, o = (para(x) for x in parca.split(":", 1))
        if e is None or o is None or e < 0 or o < 0:
            return None
        l.append((e, o))
    return l if l and [e for e, _ in l] == sorted({e for e, _ in l}) else None


def _oran(pay, payda) -> Decimal | None:
    return yuvarla(Decimal(pay) * 100 / Decimal(payda)) if payda else None


def degerlendir(satislar: list[dict], hedefler: dict | None = None, kademeler: list | None = None, esik: Decimal = Decimal(80)) -> dict:
    if not satislar:
        raise ValueError("Değerlendirilecek satış kaydı yok.")
    hedefler, uy, grup = hedefler or {}, [], {}
    for s in satislar:
        grup.setdefault(katla(s["personel"]), []).append(s)
    ciro, fis, adet = (sum((s[a] for s in satislar), SIFIR) for a in ("ciro", "fis", "adet"))
    m_sepet, m_upt = (ciro / fis if fis else None), (adet / fis if fis else None)
    personel = []
    for anahtar, g in grup.items():
        p = {"personel": g[0]["personel"], **{a: sum((s[a] for s in g), SIFIR) for a in ("ciro", "fis", "adet")}}
        gunler = {s["tarih"] for s in g if s["tarih"]}
        p["gun"] = len(gunler) or None
        p["gunluk"] = yuvarla(p["ciro"] / p["gun"]) if p["gun"] else None
        p["sepet"] = yuvarla(p["ciro"] / p["fis"]) if p["fis"] else None
        p["upt"] = yuvarla(p["adet"] / p["fis"]) if p["fis"] else None
        p["pay"] = _oran(p["ciro"], ciro)
        p["sepet_endeks"] = yuvarla(p["ciro"] / p["fis"] / m_sepet * 100) if p["fis"] and m_sepet else None
        p["upt_endeks"] = yuvarla(p["adet"] / p["fis"] / m_upt * 100) if p["fis"] and m_upt else None
        p["hedef"] = hedefler[anahtar][1] if anahtar in hedefler else None
        p["gerceklesme"] = _oran(p["ciro"], p["hedef"]) if p["hedef"] else None
        p["prim_oran"] = p["prim"] = None
        if p["gerceklesme"] is not None and kademeler:
            p["prim_oran"] = next((o for e, o in reversed(kademeler) if p["gerceklesme"] >= e), SIFIR)
            p["prim"] = yuvarla(p["ciro"] * p["prim_oran"] / 100)
        if hedefler and p["hedef"] is None:
            uy.append(uyari("Bilgi", "Hedefi tanımlı değil", p["personel"], f"Ciro {p['ciro']:,.2f}; hedef dosyasında bu personel yok, prim hesaplanmadı."))
        elif p["gerceklesme"] is not None and p["gerceklesme"] < esik:
            uy.append(uyari("Bilgi", "Hedefin gerisinde", p["personel"], f"Ciro {p['ciro']:,.2f}, hedef {p['hedef']:,.2f} (%{p['gerceklesme']}; eşik %{esik})."))
        personel.append(p)
    for anahtar, (ad, hedef) in hedefler.items():
        if anahtar not in grup:
            uy.append(uyari("Orta", "Hedefi var, satışı yok", ad, f"Hedef {hedef:,.2f} ama satış dosyasında kaydı yok. Ad yazımını ya da izin / rapor durumunu kontrol edin."))
    personel.sort(key=lambda p: (-p["ciro"], katla(p["personel"])))
    for n, p in enumerate(personel, 1):
        p["sira"] = n
    return {"personel": personel, "magaza": {"ciro": ciro, "fis": fis, "adet": adet, "sepet": yuvarla(m_sepet) if m_sepet else None, "upt": yuvarla(m_upt) if m_upt else None, "kisi": len(personel),
                                              "prim": sum((p["prim"] for p in personel if p["prim"] is not None), SIFIR) if kademeler and hedefler else None}, "kademeler": kademeler or [], "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Personel"
    _baslik(ws, ["Sıra", "Personel", "Ciro", "Ciro Payı %", "Fiş", "Satılan Adet", "Sepet Ortalaması", "Sepet, Mağazaya Göre %", "Fiş Başına Adet", "Adet, Mağazaya Göre %", "Çalışılan Gün", "Günlük Ortalama Ciro"],
            (6, 24, 16, 10, 8, 10, 13, 12, 11, 12, 9, 15), "C2")
    for p in r["personel"]:
        _satir(ws, [p["sira"], p["personel"], p["ciro"], p["pay"], p["fis"], p["adet"], p["sepet"], p["sepet_endeks"], p["upt"], p["upt_endeks"], p["gun"], p["gunluk"]], para_s=(3, 4, 7, 8, 9, 10, 12),
               dolgu={8: None if p["sepet_endeks"] is None else YESIL if p["sepet_endeks"] >= 100 else SARI, 10: None if p["upt_endeks"] is None else YESIL if p["upt_endeks"] >= 100 else SARI})
    m = r["magaza"]
    n = _satir(ws, [None, "MAĞAZA", m["ciro"], Decimal(100), m["fis"], m["adet"], m["sepet"], Decimal(100), m["upt"], Decimal(100), None, None], para_s=(3, 4, 7, 8, 9, 10))
    for h in ws[n]:
        h.font = Font(bold=True)
    if any(p["hedef"] is not None for p in r["personel"]):
        ws = wb.create_sheet("Prim Tablosu")
        _baslik(ws, ["Personel", "Ciro", "Hedef Ciro", "Gerçekleşme %", "Prim Oranı %", "Prim Tutarı", "Onay (siz doldurun)"], (24, 16, 16, 12, 11, 15, 20))
        for p in sorted(r["personel"], key=lambda p: (p["gerceklesme"] is None, -(p["gerceklesme"] or 0), katla(p["personel"]))):
            _satir(ws, [p["personel"], p["ciro"], p["hedef"], p["gerceklesme"], p["prim_oran"], p["prim"], None], para_s=(2, 3, 4, 5, 6),
                   dolgu={4: None if p["gerceklesme"] is None else YESIL if p["gerceklesme"] >= 100 else SARI})
        if m["prim"] is not None:
            n = _satir(ws, ["TOPLAM", m["ciro"], None, None, None, m["prim"], None], para_s=(2, 6))
            for h in ws[n]:
                h.font = Font(bold=True)
        ws.append([])
        ws.append(["Kademeler: " + ("; ".join(f"gerçekleşme ≥ %{e} → cironun %{o}'i" for e, o in r["kademeler"]) if r["kademeler"] else "verilmedi, prim hesaplanmadı")])
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(s_yol: Path, cikti: Path, h_yol: Path | None = None, kademeler: list | None = None, esik: Decimal = Decimal(80)) -> dict:
    satislar, uy1 = oku_satis(s_yol)
    hedefler, uy2 = oku_hedef(h_yol) if h_yol else ({}, [])
    r = degerlendir(satislar, hedefler, kademeler, esik)
    r["uyarilar"] = sirala(uy1 + uy2 + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Satış danışmanlarının ciro, adet ve sepet ortalamalarını karşılaştırır, prim tablosunu hazırlar.")
    p.add_argument("--satislar", type=Path, help="Tarih, Personel, Ciro, Fiş Adedi, Satılan Adet")
    p.add_argument("--hedefler", type=Path, help="İsteğe bağlı: Personel, Hedef Ciro")
    p.add_argument("--kademeler", help="Prim kademeleri, gerçekleşme%%:prim%% (ör. 90:1;100:2;120:3); verilmezse prim hesaplanmaz")
    p.add_argument("--esik", default="80", help="Hedef gerçekleşmesi bu yüzdenin altındaysa bildirilir (varsayılan 80)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "personel_satis_performansi.xlsx")
    a = p.parse_args(argv)
    ornek = a.satislar is None
    s_yol, h_yol = (ORNEK / "personel_satislari.csv", a.hedefler or ORNEK / "hedefler.csv") if ornek else (a.satislar, a.hedefler)
    kademe_metni = a.kademeler or ("90:1;100:2;120:3" if ornek else None)
    kademeler, esik = kademe_oku(kademe_metni) if kademe_metni else None, para(a.esik)
    if (kademe_metni and kademeler is None) or esik is None or esik < 0:
        print("[X] --kademeler artan eşiklerle gerçekleşme:prim biçiminde olmalı (ör. 90:1;100:2;120:3); --esik eksi olamaz")
        return 2
    for y in (s_yol, h_yol):
        if y is not None and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(s_yol, a.cikti, h_yol, kademeler, esik)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    m = r["magaza"]
    print(f"[OK] {m['kisi']} personel · ciro {m['ciro']:,.2f} · fiş {m['fis']:,.0f}" + (f" · sepet {m['sepet']:,.2f}" if m["sepet"] is not None else "") + (f" · toplam prim {m['prim']:,.2f}" if m["prim"] is not None else ""))
    for p in r["personel"][:3]:
        print(f"     {p['sira']}. {p['personel']} · ciro {p['ciro']:,.2f} (%{p['pay']})" + (f" · hedef %{p['gerceklesme']}" if p["gerceklesme"] is not None else ""))
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
