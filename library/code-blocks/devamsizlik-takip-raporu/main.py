"""
Devamsızlık Takip Raporu — Workers / Workless kod bloğu
Kayıt Kabul ve Öğrenci İşleri › Öğrenci İşleri Sorumlusu
Yoklama kayıtlarından öğrenci bazında özürlü / özürsüz devamsızlığı hesaplar, sınıra yaklaşan ve aşan öğrencileri, veli bildirimi gerekenleri listeler.
  Gün          = tam gün 1, yarım gün 0,5 (aynı gün için birden fazla kayıt toplanır, en fazla 1)
  Lise sınırı  = özürsüz 10 gün, toplam 30 gün (Ortaöğretim Kurumları Yönetmeliği m.36; birinci derece yakın kaybında toplam 40 — "Toplam Sınır" sütunuyla)
  İlkokul / ortaokul = kodda gün sınırı yok; --ilk-toplam / --ilk-ozursuz verilirse o eşiklerle izlenir
  Yaklaşıyor   = sınırın --esik (%80) ve üstü
  Veli bildirimi = art arda 2 okul günü (tam gün) özürsüz devamsızlık; özürsüzün 5. gününde veli daveti (lise)
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
YOKLAMA = {"no": ("ogrenci no", "okul no", "no", "numara"), "ad": ("ad soyad", "ogrenci", "adi soyadi", "ogrenci adi"), "sinif": ("sinif", "sube", "sinif sube"),
           "kademe": ("kademe", "okul turu", "seviye"), "tarih": ("tarih", "gun"), "sure": ("sure", "devamsizlik", "gun sayisi", "tam yarim"),
           "tur": ("tur", "ozur", "ozur durumu", "devamsizlik turu"), "sinir": ("toplam sinir", "ozel sinir", "ek hak")}
SIFIR = Decimal(0)


def _sure(x) -> Decimal | None:
    k = katla(x)
    if not k or k in ("tam", "tam gun", "1"):
        return Decimal(1)
    if k in ("yarim", "yarim gun", "0 5", "0,5", "05"):
        return Decimal("0.5")
    p = para(x)
    return p if p is not None and p in (Decimal("0.5"), Decimal(1)) else None


def _tur(x) -> str | None:
    k = katla(x)
    if k in ("ozurlu", "raporlu", "izinli", "ozur", "e"):
        return "Özürlü"
    if not k or k in ("ozursuz", "raporsuz", "h"):
        return "Özürsüz"
    return None


def _kademe(x) -> str | None:
    k = katla(x)
    if "lise" in k or "orta ogretim" in k or "ortaogretim" in k:
        return "Lise"
    if "ortaokul" in k or "imam hatip ortaokulu" in k:
        return "Ortaokul"
    if "ilkokul" in k:
        return "İlkokul"
    return None if not k else "Diğer"


def oku(yol: Path, varsayilan: str = "Lise") -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, YOKLAMA, ("no", "tarih")):
        no = metin(r["no"])
        if not no:
            continue
        g, s, t = tarih(r["tarih"]), _sure(r["sure"]), _tur(r["tur"])
        sn = para(r["sinir"]) if metin(r["sinir"]) else None
        if g is None or s is None or t is None or (metin(r["sinir"]) and (sn is None or sn <= 0)):
            uy.append(uyari("Yüksek", "Satır okunamadı", f"{no} {metin(r['ad'])}".strip(), f"Satır {r['_satir']}: tarih, süre ('{metin(r['sure'])}'), tür ('{metin(r['tur'])}') ya da sınır okunamadı."))
            continue
        l.append({"no": no, "ad": metin(r["ad"]), "sinif": metin(r["sinif"]), "kademe": _kademe(r["kademe"]) or varsayilan, "tarih": g, "sure": s, "tur": t, "sinir": sn})
    return l, uy


def _g(x) -> str:
    return f"{x:.1f}".rstrip("0").rstrip(".").replace(".", ",")


def _onceki_okul_gunu(g: date) -> date:
    g -= timedelta(days=1)
    while g.weekday() >= 5:
        g -= timedelta(days=1)
    return g


def analiz(kayit: list[dict], esik: Decimal = Decimal(80), ilk_toplam: Decimal | None = None, ilk_ozursuz: Decimal | None = None) -> dict:
    if not kayit:
        raise ValueError("Okunabilen yoklama kaydı yok.")
    uy, ogr = [], {}
    for k in kayit:
        o = ogr.setdefault(k["no"], {"no": k["no"], "ad": k["ad"], "sinif": k["sinif"], "kademe": k["kademe"], "sinir": None, "gunler": {}})
        o["ad"], o["sinif"] = o["ad"] or k["ad"], o["sinif"] or k["sinif"]
        o["sinir"] = k["sinir"] or o["sinir"]
        gun = o["gunler"].setdefault(k["tarih"], {"Özürlü": SIFIR, "Özürsüz": SIFIR})
        gun[k["tur"]] += k["sure"]
    satirlar = []
    for o in ogr.values():
        for g, d in o["gunler"].items():
            top = d["Özürlü"] + d["Özürsüz"]
            if top > 1:
                fazla = top - 1
                d["Özürlü"] = max(d["Özürlü"] - fazla, SIFIR)
                if d["Özürlü"] + d["Özürsüz"] > 1:
                    d["Özürsüz"] = Decimal(1) - d["Özürlü"]
                uy.append(uyari("Bilgi", "Aynı gün fazla kayıt", f"{o['no']} {o['ad']}".strip(), f"{g:%d.%m.%Y} için toplam {_g(top)} gün yazılmış; 1 gün sayıldı."))
        oz = sum((d["Özürsüz"] for d in o["gunler"].values()), SIFIR)
        ol = sum((d["Özürlü"] for d in o["gunler"].values()), SIFIR)
        top = oz + ol
        if o["kademe"] == "Lise":
            s_oz, s_top = Decimal(10), o["sinir"] or Decimal(30)
        else:
            s_oz, s_top = ilk_ozursuz, o["sinir"] or ilk_toplam
        kim = f"{o['no']} {o['ad']}".strip()
        durum = "Normal"
        if (s_oz is not None and oz > s_oz) or (s_top is not None and top > s_top):
            durum = "Sınır aşıldı"
            uy.append(uyari("Yüksek", "Devamsızlık sınırı aşıldı", kim, f"{o['sinif']} · özürsüz {_g(oz)}" + (f"/{_g(s_oz)}" if s_oz is not None else "") + f", toplam {_g(top)}" + (f"/{_g(s_top)}" if s_top is not None else "")
                            + (" gün. Lisede sınırı aşan öğrenci ders puanlarına bakılmaksızın başarısız sayılır; veliye yazılı bildirim yapılır." if o["kademe"] == "Lise" else " gün.")))
        elif (s_oz is not None and oz * 100 >= s_oz * esik) or (s_top is not None and top * 100 >= s_top * esik):
            durum = "Sınıra yaklaştı"
            uy.append(uyari("Orta", "Sınıra yaklaştı", kim, f"{o['sinif']} · özürsüz {_g(oz)}, toplam {_g(top)} gün" + (f" (kalan özürsüz {_g(s_oz - oz)}, toplam {_g(s_top - top)})" if s_oz is not None and s_top is not None else "") + "."))
        tam_ozursuz = sorted(g for g, d in o["gunler"].items() if d["Özürsüz"] >= 1)
        ardisik = [g for g in tam_ozursuz if _onceki_okul_gunu(g) in set(tam_ozursuz)]
        for g in ardisik:
            uy.append(uyari("Orta", "Art arda 2 gün özürsüz", kim, f"{_onceki_okul_gunu(g):%d.%m.%Y} ve {g:%d.%m.%Y} günleri özürsüz devamsız; durum veliye bildirilmeli."))
        besinci = None
        if o["kademe"] == "Lise" and oz >= 5:
            biriken = SIFIR
            for g in sorted(o["gunler"]):
                biriken += o["gunler"][g]["Özürsüz"]
                if biriken >= 5:
                    besinci = g
                    break
            uy.append(uyari("Bilgi", "Özürsüz 5. gün", kim, f"Özürsüz devamsızlık {besinci:%d.%m.%Y} tarihinde 5 güne ulaştı; sınıf tekrarı riski için veli okula davet edilir."))
        satirlar.append({"no": o["no"], "ad": o["ad"], "sinif": o["sinif"], "kademe": o["kademe"], "ozursuz": oz, "ozurlu": ol, "toplam": top, "s_oz": s_oz, "s_top": s_top,
                         "kalan_oz": s_oz - oz if s_oz is not None else None, "kalan_top": s_top - top if s_top is not None else None, "durum": durum,
                         "ardisik": len(ardisik), "son": max(o["gunler"]), "besinci": besinci})
    sira = {"Sınır aşıldı": 0, "Sınıra yaklaştı": 1, "Normal": 2}
    satirlar.sort(key=lambda s: (sira[s["durum"]], -s["toplam"], s["sinif"], s["no"]))
    siniflar = {}
    for s in satirlar:
        c = siniflar.setdefault(s["sinif"] or "-", {"sinif": s["sinif"] or "-", "ogrenci": 0, "ozursuz": SIFIR, "ozurlu": SIFIR, "asan": 0, "yakin": 0})
        c["ogrenci"] += 1
        c["ozursuz"] += s["ozursuz"]
        c["ozurlu"] += s["ozurlu"]
        c["asan"] += s["durum"] == "Sınır aşıldı"
        c["yakin"] += s["durum"] == "Sınıra yaklaştı"
    return {"ogrenciler": satirlar, "siniflar": sorted(siniflar.values(), key=lambda c: c["sinif"]), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Öğrenciler"
    _baslik(ws, ["Öğrenci No", "Ad Soyad", "Sınıf", "Kademe", "Özürsüz", "Özürlü", "Toplam", "Özürsüz Sınırı", "Toplam Sınırı", "Kalan Özürsüz", "Kalan Toplam", "Durum",
                 "Art Arda Özürsüz", "Özürsüz 5. Gün", "Son Devamsızlık"], (10, 22, 8, 9, 8, 8, 8, 9, 9, 9, 9, 15, 9, 11, 11), "C2")
    renk = {"Sınır aşıldı": KIRMIZI, "Sınıra yaklaştı": SARI}
    for s in r["ogrenciler"]:
        _satir(ws, [s["no"], s["ad"], s["sinif"], s["kademe"], s["ozursuz"], s["ozurlu"], s["toplam"], s["s_oz"], s["s_top"], s["kalan_oz"], s["kalan_top"], s["durum"], s["ardisik"] or None,
                    s["besinci"], s["son"]], tarih_s=(14, 15), dolgu={12: renk.get(s["durum"])})
    ws = wb.create_sheet("Sınıflar")
    _baslik(ws, ["Sınıf", "Devamsız Öğrenci", "Özürsüz Gün", "Özürlü Gün", "Sınırı Aşan", "Sınıra Yaklaşan"], (10, 10, 10, 10, 10, 10))
    for c in r["siniflar"]:
        _satir(ws, [c["sinif"], c["ogrenci"], c["ozursuz"], c["ozurlu"], c["asan"], c["yakin"]])
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, esik: Decimal = Decimal(80), ilk_toplam: Decimal | None = None, ilk_ozursuz: Decimal | None = None, varsayilan: str = "Lise") -> dict:
    kayit, uy = oku(yol, varsayilan)
    r = analiz(kayit, esik, ilk_toplam, ilk_ozursuz)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Öğrenci bazında özürlü / özürsüz devamsızlığı hesaplar, sınıra yaklaşanları ve veli bildirimi gerekenleri listeler.")
    p.add_argument("--yoklama", type=Path, help="Öğrenci No, Ad Soyad, Sınıf, Kademe (İlkokul / Ortaokul / Lise), Tarih, Süre (Tam / Yarım), Tür (Özürlü / Özürsüz), Toplam Sınır (isteğe bağlı)")
    p.add_argument("--esik", default="80", help="Sınırın yüzde kaçında 'yaklaştı' uyarısı verilir (varsayılan 80)")
    p.add_argument("--ilk-toplam", help="İlkokul / ortaokul için izlenecek toplam gün eşiği (varsayılan yok)")
    p.add_argument("--ilk-ozursuz", help="İlkokul / ortaokul için izlenecek özürsüz gün eşiği (varsayılan yok)")
    p.add_argument("--kademe", choices=("Lise", "Ortaokul", "İlkokul"), default="Lise", help="Kademe sütunu boşsa kullanılacak kademe (varsayılan Lise)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "devamsizlik_raporu.xlsx")
    a = p.parse_args(argv)
    yol = a.yoklama or ORNEK / "yoklama.csv"
    esik = para(a.esik)
    it, io = (para(a.ilk_toplam) if a.ilk_toplam else None), (para(a.ilk_ozursuz) if a.ilk_ozursuz else None)
    if esik is None or not 1 <= esik <= 100 or (a.ilk_toplam and (it is None or it <= 0)) or (a.ilk_ozursuz and (io is None or io <= 0)):
        print("[X] --esik 1-100; --ilk-toplam ve --ilk-ozursuz 0'dan büyük olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, esik, it, io, a.kademe)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    l = r["ogrenciler"]
    print(f"[OK] {len(l)} devamsız öğrenci · sınırı aşan {sum(1 for s in l if s['durum'] == 'Sınır aşıldı')} · sınıra yaklaşan {sum(1 for s in l if s['durum'] == 'Sınıra yaklaştı')} · "
          f"art arda özürsüz {sum(1 for s in l if s['ardisik'])}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
