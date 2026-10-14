import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
PZT = date(2026, 3, 2)


def k(no, gun, tur="Özürsüz", sure=1, kademe="Lise", sinir=None):
    return {"no": no, "ad": f"Öğrenci {no}", "sinif": "9-A", "kademe": kademe, "tarih": PZT + timedelta(days=gun), "sure": D(sure), "tur": tur, "sinir": None if sinir is None else D(sinir)}


def gunler(no, adet, tur, bas=0, **kw):
    """Hafta içi günlere adet kadar kayıt."""
    l, g = [], bas
    while len(l) < adet:
        if (PZT + timedelta(days=g)).weekday() < 5:
            l.append(k(no, g, tur, **kw))
        g += 1
    return l


class AnalizTesti(unittest.TestCase):
    def test_sinirlar(self):
        kayit = (gunler("1", 11, "Özürsüz") + gunler("2", 8, "Özürsüz", 0) + gunler("2", 16, "Özürlü", 30) + gunler("3", 25, "Özürlü", 0) + gunler("3", 6, "Özürsüz", 50)
                 + gunler("4", 25, "Özürlü", 0, sinir=40) + gunler("4", 6, "Özürsüz", 50, sinir=40) + gunler("5", 31, "Özürlü", kademe="Ortaokul"))
        r = main.analiz(kayit)
        s = {x["no"]: x for x in r["ogrenciler"]}
        self.assertEqual([(n, s[n]["ozursuz"], s[n]["toplam"], s[n]["durum"]) for n in "12345"],
                         [("1", D(11), D(11), "Sınır aşıldı"), ("2", D(8), D(24), "Sınıra yaklaştı"), ("3", D(6), D(31), "Sınır aşıldı"), ("4", D(6), D(31), "Normal"),
                          ("5", D(0), D(31), "Normal")])
        self.assertEqual((s["4"]["s_top"], s["4"]["kalan_top"], s["5"]["s_top"]), (D(40), D(9), None))
        self.assertEqual((s["1"]["ardisik"] > 0, s["1"]["besinci"]), (True, date(2026, 3, 6)))
        self.assertEqual(main.analiz(gunler("5", 31, "Özürlü", kademe="Ortaokul"), ilk_toplam=D(30))["ogrenciler"][0]["durum"], "Sınır aşıldı")
        self.assertEqual([c["asan"] for c in r["siniflar"]], [2])
        with self.assertRaises(ValueError):
            main.analiz([])

    def test_gun_ve_ardisik(self):
        # cuma + pazartesi art arda okul günü sayılır; yarım gün 0,5; aynı gün fazlası 1'e indirilir
        kayit = [k("1", 4), k("1", 7), k("1", 9, sure="0.5"), k("1", 9, sure="0.5", tur="Özürlü"), k("1", 10), k("1", 10, tur="Özürlü")]
        r = main.analiz(kayit)
        o = r["ogrenciler"][0]
        self.assertEqual((o["ozursuz"], o["ozurlu"], o["toplam"], o["ardisik"]), (D("3.5"), D("0.5"), D(4), 1))
        self.assertEqual(sorted(u["tur"] for u in r["uyarilar"]), ["Art arda 2 gün özürsüz", "Aynı gün fazla kayıt"])
        self.assertEqual(main._onceki_okul_gunu(date(2026, 3, 9)), date(2026, 3, 6))

    def test_donusum(self):
        self.assertEqual([main._sure(x) for x in ("", "Yarım gün", "0,5", "1", "2")], [D(1), D("0.5"), D("0.5"), D(1), None])
        self.assertEqual([main._tur(x) for x in ("", "Raporlu", "x")], ["Özürsüz", "Özürlü", None])
        self.assertEqual([main._kademe(x) for x in ("Anadolu Lisesi", "Ortaokul", "İlkokul", "", "x")], ["Lise", "Ortaokul", "İlkokul", None, "Diğer"])
        self.assertEqual(main._g(D("13.0")), "13")


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "yoklama.csv", cikti)
            l = r["ogrenciler"]
            self.assertEqual((len(l), sum(1 for s in l if s["durum"] == "Sınır aşıldı"), sum(1 for s in l if s["durum"] == "Sınıra yaklaştı")), (59, 1, 2))
            self.assertEqual(l[0]["no"], "116")
            self.assertEqual(load_workbook(cikti).sheetnames, ["Öğrenciler", "Sınıflar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "y.csv"
            y.write_text("Okul No;Öğrenci;Şube;Tarih;Süre;Özür Durumu;Ek Hak\n1;A;9-A;02.03.2026;;;\n2;B;;x;;;\n3;C;;02.03.2026;3;;\n4;D;;02.03.2026;;belki;\n5;E;;02.03.2026;;;0\n", encoding="utf-8")
            l, uy = main.oku(y, "Ortaokul")
            self.assertEqual(([(x["no"], x["kademe"], x["tur"]) for x in l], len(uy)), ([("1", "Ortaokul", "Özürsüz")], 4))
            c = str(Path(tmp) / "x.xlsx")
            self.assertEqual(main.main(["--yoklama", str(y), "--esik", "70", "--ilk-toplam", "20", "--ilk-ozursuz", "10", "--kademe", "İlkokul", "--cikti", c]), 0)
            self.assertEqual([main.main(["--yoklama", str(y), "--cikti", c] + v) for v in (["--esik", "0"], ["--ilk-toplam", "x"], ["--ilk-ozursuz", "0"])], [2] * 3)
            self.assertEqual(main.main(["--yoklama", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Okul No;Tarih\n1;x\n", encoding="utf-8")
            self.assertEqual(main.main(["--yoklama", str(y), "--cikti", c]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
