"""
Cari Hesap Ekstresi Hazırlama — Workers / Workless kod bloğu
Muhasebe › Ön Muhasebe Elemanı

Muhasebe hareketlerinden müşteri / tedarikçi bazında dönem ekstresi ve bakiye özeti hazırlar:
  - Devreden bakiye (dönem başından önceki hareketler), dönem hareketleri, yürüyen bakiye, kapanış bakiyesi (B/A).
  - Açık kalemler FIFO ile bulunur: bakiyenin yönündeki kalemler (müşteride faturalar, tedarikçide alış
    faturaları) en eskiden kapatılır; kalanların vadesi dönem sonundan önceyse "vadesi geçmiş" sayılır.
    Vade tarihi yoksa cari kartındaki vade günü (yoksa --vade) belge tarihine eklenir.
  - Kontroller: ters bakiye (alacak bakiyeli müşteri / borç bakiyeli tedarikçi → avans veya fazla ödeme),
    dönemde hareketi olmayan bakiye, aynı belge no + tutarın tekrarı, cari kartı olmayan hareket.
  - Her cari için mutabakat yazısı taslağı; istenirse her cari için ayrı ekstre dosyası (--ayri-dosya).
Rapor: bakiye özeti, ekstreler, açık kalemler, mutabakat metinleri, uyarılar. İnternete bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: 6 cari, 01.07–30.09.2026
    python main.py --hareketler hareketler.xlsx --cariler cariler.xlsx --donem 01.07.2026 30.09.2026
    python main.py --hareketler h.xlsx --donem 01.01.2026 30.09.2026 --ayri-dosya --firma "Örnek AŞ"
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
K2 = Decimal("0.01")
SIFIR = Decimal(0)

HAREKET_SUTUNLARI = {"tarih": ("tarih", "belge tarihi", "islem tarihi"), "kod": ("cari kod", "cari kodu", "hesap kodu", "cari hesap kodu"),
                     "unvan": ("cari unvan", "unvan", "cari adi"), "tur": ("belge turu", "islem turu", "fis turu"), "no": ("belge no", "evrak no", "fatura no"),
                     "aciklama": ("aciklama",), "borc": ("borc", "borc tutari"), "alacak": ("alacak", "alacak tutari"), "vade": ("vade tarihi", "vade")}
CARI_SUTUNLARI = {"kod": ("cari kod", "cari kodu", "hesap kodu"), "unvan": ("unvan", "cari unvan"), "tur": ("tur", "cari turu", "tip"),
                  "vade": ("vade gun", "vade", "odeme vadesi")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tl(x: Decimal | None) -> str:
    return "—" if x is None else f"{x.quantize(K2, ROUND_HALF_UP):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def ba(x: Decimal) -> str:
    return "B" if x > 0 else "A" if x < 0 else ""


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Hareket:
    tarih: date
    tur: str
    no: str
    aciklama: str
    borc: Decimal
    alacak: Decimal
    vade: date | None
    sira: int

    @property
    def net(self) -> Decimal:
        return self.borc - self.alacak


@dataclass
class Cari:
    kod: str
    unvan: str
    tur: str = ""
    vade_gun: int | None = None
    hareketler: list = field(default_factory=list)
    devir: Decimal = SIFIR
    donem: list = field(default_factory=list)
    acik: list = field(default_factory=list)          # (hareket, açık tutar, vade)

    @property
    def borc(self) -> Decimal:
        return sum((h.borc for h in self.donem), SIFIR)

    @property
    def alacak(self) -> Decimal:
        return sum((h.alacak for h in self.donem), SIFIR)

    @property
    def bakiye(self) -> Decimal:
        return self.devir + self.borc - self.alacak

    @property
    def musteri(self) -> bool | None:
        k = katla(self.tur)
        if k.startswith(("musteri", "alici")):
            return True
        if k.startswith(("tedarikci", "satici")):
            return False
        return True if self.kod.startswith("12") else False if self.kod.startswith("32") else None


def oku(hareket_yolu: Path, cari_yolu: Path | None, varsayilan_vade: int) -> tuple[dict[str, Cari], list[dict]]:
    uy = []
    cariler: dict[str, Cari] = {}
    if cari_yolu:
        for r in kayitlar(cari_yolu, CARI_SUTUNLARI, ("kod",)):
            if metin(r.get("kod")):
                v = para(r.get("vade"))
                cariler[metin(r["kod"])] = Cari(metin(r["kod"]), metin(r.get("unvan")), metin(r.get("tur")), int(v) if v is not None else None)
    for r in kayitlar(hareket_yolu, HAREKET_SUTUNLARI, ("tarih", "kod")):
        kod, t = metin(r.get("kod")), tarih(r.get("tarih"))
        if not kod:
            continue
        if not t:
            uy.append({"onem": "Orta", "tur": "Okunamayan satır", "cari": kod, "aciklama": f"Satır {r['_satir']}: tarih okunamadı"})
            continue
        b, a = para(r.get("borc")) or SIFIR, para(r.get("alacak")) or SIFIR
        if kod not in cariler:
            cariler[kod] = Cari(kod, metin(r.get("unvan")) or kod)
            if cari_yolu:
                uy.append({"onem": "Bilgi", "tur": "Cari kartı yok", "cari": kod, "aciklama": "Hareket var, cari listesinde kayıt yok"})
        c = cariler[kod]
        if not c.unvan and metin(r.get("unvan")):
            c.unvan = metin(r.get("unvan"))
        c.hareketler.append(Hareket(t, metin(r.get("tur")), metin(r.get("no")), metin(r.get("aciklama")), b, a, tarih(r.get("vade")), r["_satir"]))
    for c in cariler.values():
        c.hareketler.sort(key=lambda h: (h.tarih, h.sira))
        gun = c.vade_gun if c.vade_gun is not None else varsayilan_vade
        for h in c.hareketler:
            if h.vade is None and gun is not None and dogal_kalem(c, h):
                h.vade = h.tarih + timedelta(days=gun)
    return cariler, uy


ODEME_TURLERI = ("tahsil", "odeme", "iade", "cek", "senet", "havale", "eft", "nakit", "virman", "mahsup", "kredi karti")


def dogal_kalem(c: Cari, h: Hareket) -> bool:
    """Vadesi olan kalem mi: müşteride borç, tedarikçide alacak kaydı ve ödeme / iade türünde değil."""
    if any(k in katla(h.tur) for k in ODEME_TURLERI):
        return False
    m = c.musteri
    return (h.borc > 0) if m is True else (h.alacak > 0) if m is False else True


# ----------------------------------------------------------------------------
# Hesap
# ----------------------------------------------------------------------------

def fifo_acik(c: Cari, son: date) -> list[tuple[Hareket, Decimal]]:
    """Kapanış bakiyesinin yönündeki kalemleri en eskiden kapatır; açık kalanları döndürür."""
    hs = [h for h in c.hareketler if h.tarih <= son]
    bakiye = sum((h.net for h in hs), SIFIR)
    if bakiye == 0:
        return []
    isaret = 1 if bakiye > 0 else -1
    kalemler = [(h, h.net * isaret) for h in hs if h.net * isaret > 0]
    kapatici = sum((-h.net * isaret for h in hs if h.net * isaret < 0), SIFIR)
    acik = []
    for h, t in kalemler:
        if kapatici >= t:
            kapatici -= t
            continue
        acik.append((h, t - kapatici))
        kapatici = SIFIR
    return acik


def analiz_et(cariler: dict[str, Cari], bas: date, son: date) -> list[dict]:
    uyarilar = []
    for c in cariler.values():
        c.devir = sum((h.net for h in c.hareketler if h.tarih < bas), SIFIR)
        c.donem = [h for h in c.hareketler if bas <= h.tarih <= son]
        c.acik = fifo_acik(c, son)
        tekrar = Counter((h.no, h.net) for h in c.hareketler if h.no)
        for (no, net), n in tekrar.items():
            if n > 1:
                uyarilar.append({"onem": "Yüksek", "tur": "Mükerrer kayıt", "cari": c.kod, "aciklama": f"{no} belge no ve {tl(abs(net))} TL tutarla {n} kez kayıtlı"})
        m = c.musteri
        if m is True and c.bakiye < 0:
            uyarilar.append({"onem": "Orta", "tur": "Ters bakiye", "cari": c.kod, "aciklama": f"Müşteri alacak bakiyesi veriyor ({tl(-c.bakiye)} TL): "
                             "avans veya fazla tahsilat; iade ya da mahsup değerlendirin"})
        if m is False and c.bakiye > 0:
            uyarilar.append({"onem": "Orta", "tur": "Ters bakiye", "cari": c.kod, "aciklama": f"Tedarikçi borç bakiyesi veriyor ({tl(c.bakiye)} TL): "
                             "avans veya fazla ödeme"})
        if not c.donem and c.bakiye:
            son_h = max((h.tarih for h in c.hareketler if h.tarih < bas), default=None)
            uyarilar.append({"onem": "Orta", "tur": "Hareketsiz bakiye", "cari": c.kod, "aciklama": f"Dönemde hareket yok, bakiye {tl(abs(c.bakiye))} "
                             f"{ba(c.bakiye)}" + (f"; son hareket {son_h:%d.%m.%Y}" if son_h else "") + ". Tahsil edilemeyen / unutulan kalem olabilir"})
        gecmis = sum((t for h, t in c.acik if h.vade and h.vade < son and dogal_kalem(c, h)), SIFIR)
        if gecmis:
            en_eski = min(h.vade for h, t in c.acik if h.vade and h.vade < son and dogal_kalem(c, h))
            uyarilar.append({"onem": "Bilgi", "tur": "Vadesi geçmiş", "cari": c.kod, "aciklama": f"{tl(gecmis)} TL vadesi geçmiş; en eski vade "
                             f"{en_eski:%d.%m.%Y} ({(son - en_eski).days} gün)"})
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    uyarilar.sort(key=lambda u: (sira[u["onem"]], u["tur"], u["cari"]))
    return uyarilar


def mutabakat_metni(c: Cari, son: date, firma: str, sure: int) -> str:
    if c.bakiye == 0:
        durum = "bakiye vermemektedir"
    else:
        taraf = "şirketimize borç" if c.bakiye > 0 else "şirketimizden alacak"
        durum = f"{tl(abs(c.bakiye))} TL {taraf} bakiyesi vermektedir"
    return (f"Sayın {c.unvan},\n\n{son:%d.%m.%Y} tarihi itibarıyla kayıtlarımızda cari hesabınız {durum}. Hesap ekstresi ektedir.\n\n"
            f"Kayıtlarınızla mutabık olup olmadığınızı {sure} gün içinde bildirmenizi; mutabık değilseniz farklılık gösteren kalemleri "
            f"kendi ekstrenizle birlikte iletmenizi rica ederiz.\n\nSaygılarımızla,\n{firma or '[Firma adı]'}")


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
GRUP = PatternFill("solid", fgColor="EEEEEE")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KONTROL = PatternFill("solid", fgColor="FFF4CE")
UST = Alignment(vertical="top", wrap_text=True)
PF = "#,##0.00"


def _baslik(ws, basliklar, genislik=()):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def ekstre_yaz(ws, c: Cari, bas: date, son: date) -> None:
    ws.append([f"{c.kod} · {c.unvan}", "", "", "", "", "", "", f"{bas:%d.%m.%Y} – {son:%d.%m.%Y}"])
    for h in ws[ws.max_row]:
        h.fill, h.font = GRUP, Font(bold=True)
    ws.append([bas - timedelta(days=1), "Devir", "", "Devreden bakiye", None, None, None, float(abs(c.devir)), ba(c.devir)])
    ws.cell(ws.max_row, 1).number_format = "DD.MM.YYYY"
    ws.cell(ws.max_row, 8).number_format = PF
    y = c.devir
    for h in c.donem:
        y += h.net
        ws.append([h.tarih, h.tur, h.no, h.aciklama, h.vade, float(h.borc) or None, float(h.alacak) or None, float(abs(y)), ba(y)])
        r = ws.max_row
        ws.cell(r, 1).number_format = ws.cell(r, 5).number_format = "DD.MM.YYYY"
        for j in (6, 7, 8):
            ws.cell(r, j).number_format = PF
    ws.append(["", "", "", "Dönem toplamı / kapanış", None, float(c.borc), float(c.alacak), float(abs(c.bakiye)), ba(c.bakiye)])
    for j in (6, 7, 8):
        ws.cell(ws.max_row, j).number_format = PF
        ws.cell(ws.max_row, j).font = Font(bold=True)
    ws.append([])


def ekstre_basligi(ws):
    _baslik(ws, ["Tarih", "Belge Türü", "Belge No", "Açıklama", "Vade", "Borç", "Alacak", "Bakiye", "B/A"], (11, 11, 19, 30, 11, 14, 14, 14, 5))


def rapor_yaz(cikti: Path, s: dict) -> None:
    bas, son = s["donem"]
    cariler = s["cariler"]
    wb = Workbook()
    oz = wb.active
    oz.title = "Bakiye Özeti"
    _baslik(oz, ["Cari Kod", "Unvan", "Tür", "Devir", "Devir B/A", "Dönem Borç", "Dönem Alacak", "Bakiye", "B/A", "Açık Kalem", "Vadesi Geçmiş",
                 "En Eski Vade", "Mutabakat Durumu"], (11, 30, 10, 14, 6, 14, 14, 14, 5, 9, 14, 11, 18))
    for c in cariler:
        gecmis = [(h, t) for h, t in c.acik if h.vade and h.vade < son and dogal_kalem(c, h)]
        oz.append([c.kod, c.unvan, c.tur, float(abs(c.devir)), ba(c.devir), float(c.borc), float(c.alacak), float(abs(c.bakiye)), ba(c.bakiye), len(c.acik),
                   float(sum((t for _, t in gecmis), SIFIR)), min((h.vade for h, _ in gecmis), default=None), ""])
        r = oz.max_row
        for j in (4, 6, 7, 8, 11):
            oz.cell(r, j).number_format = PF
        oz.cell(r, 12).number_format = "DD.MM.YYYY"
        if gecmis:
            oz.cell(r, 11).fill = PatternFill("solid", fgColor="FDE2E1")
        oz.cell(r, 13).fill = KONTROL
    oz.freeze_panes = "A2"
    oz.auto_filter.ref = f"A1:M{oz.max_row}"

    ek = wb.create_sheet("Ekstreler")
    ekstre_basligi(ek)
    for c in cariler:
        ekstre_yaz(ek, c, bas, son)

    ak = wb.create_sheet("Açık Kalemler")
    _baslik(ak, ["Cari Kod", "Unvan", "Belge Tarihi", "Belge No", "Açıklama", "Vade", "Açık Tutar", "Yön", "Gecikme (gün)"], (11, 30, 11, 19, 26, 11, 14, 5, 10))
    for c in cariler:
        for h, t in c.acik:
            gec = (son - h.vade).days if h.vade and h.vade < son and dogal_kalem(c, h) else 0
            ak.append([c.kod, c.unvan, h.tarih, h.no, h.aciklama, h.vade, float(t), ba(c.bakiye), gec])
            r = ak.max_row
            ak.cell(r, 3).number_format = ak.cell(r, 6).number_format = "DD.MM.YYYY"
            ak.cell(r, 7).number_format = PF
            if gec:
                ak.cell(r, 9).fill = PatternFill("solid", fgColor="FDE2E1")
    ak.freeze_panes = "A2"

    mt = wb.create_sheet("Mutabakat Metinleri")
    _baslik(mt, ["Cari Kod", "Unvan", "Metin", "Gönderildi", "Cevap"], (11, 30, 90, 11, 14))
    for c in cariler:
        mt.append([c.kod, c.unvan, mutabakat_metni(c, son, s["firma"], s["sure"]), "", ""])
        mt.cell(mt.max_row, 3).alignment = UST
        mt.cell(mt.max_row, 4).fill = mt.cell(mt.max_row, 5).fill = KONTROL

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Cari", "Açıklama"], (9, 18, 11, 90))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["cari"], u["aciklama"]])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def ayri_dosyalar(klasor: Path, s: dict) -> list[Path]:
    bas, son = s["donem"]
    klasor.mkdir(parents=True, exist_ok=True)
    yollar = []
    for c in s["cariler"]:
        wb = Workbook()
        ws = wb.active
        ws.title = "Ekstre"
        ws.append([s["firma"] or "", "", "", "Cari Hesap Ekstresi"])
        ws.cell(1, 1).font = ws.cell(1, 4).font = Font(bold=True, size=13)
        ws.append([])
        ekstre_basligi(ws)
        ekstre_yaz(ws, c, bas, son)
        ws.append([mutabakat_metni(c, son, s["firma"], s["sure"])])
        ws.cell(ws.max_row, 1).alignment = UST
        ad = re.sub(r"[^\w.-]+", "_", f"{c.kod}_{c.unvan}")[:80] + ".xlsx"
        wb.save(klasor / ad)
        yollar.append(klasor / ad)
    return yollar


def calistir(hareket_yolu: Path, cikti: Path, donem: tuple[date, date], cari_yolu: Path | None = None, *, vade: int | None = None, firma: str = "",
             sure: int = 15, ayri: bool = False) -> dict:
    cariler, uyarilar = oku(hareket_yolu, cari_yolu, vade)
    if not cariler:
        raise ValueError(f"{hareket_yolu.name}: hareket bulunamadı")
    uyarilar += analiz_et(cariler, *donem)
    liste = [c for c in sorted(cariler.values(), key=lambda c: c.kod) if c.donem or c.bakiye or c.devir]
    s = {"cariler": liste, "uyarilar": uyarilar, "donem": donem, "firma": firma, "sure": sure}
    rapor_yaz(cikti, s)
    s["dosyalar"] = ayri_dosyalar(cikti.parent / (cikti.stem + "_cari"), s) if ayri else []
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Muhasebe hareketlerinden cari bazında dönem ekstresi, bakiye özeti ve mutabakat metni hazırlar.")
    p.add_argument("--hareketler", type=Path, default=ORNEK / "hareketler.csv", help="Tarih, Cari Kod, Belge Türü, Belge No, Açıklama, Borç, Alacak, Vade Tarihi")
    p.add_argument("--cariler", type=Path, help="İsteğe bağlı: Cari Kod, Unvan, Tür (Müşteri/Tedarikçi), Vade (gün)")
    p.add_argument("--donem", nargs=2, metavar=("BAŞLANGIÇ", "BİTİŞ"), help="Ekstre dönemi GG.AA.YYYY GG.AA.YYYY")
    p.add_argument("--vade", type=int, help="Vade tarihi ve cari vadesi yoksa kullanılacak gün")
    p.add_argument("--firma", default="", help="Mutabakat metninde imza olarak yazılacak firma adı")
    p.add_argument("--cevap-suresi", type=int, default=15, help="Mutabakat metnindeki cevap süresi, gün (varsayılan 15)")
    p.add_argument("--ayri-dosya", action="store_true", help="Her cari için ayrı ekstre dosyası üret")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "cari_ekstreler.xlsx")
    a = p.parse_args(argv)
    ornek = a.hareketler == ORNEK / "hareketler.csv"
    cari = a.cariler or (ORNEK / "cariler.csv" if ornek else None)
    for y in (a.hareketler, cari):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    if a.donem:
        donem = (tarih(a.donem[0]), tarih(a.donem[1]))
        if None in donem or donem[0] > donem[1]:
            print("[X] --donem iki geçerli tarih olmalı: GG.AA.YYYY GG.AA.YYYY")
            return 2
    elif ornek:
        donem = (date(2026, 7, 1), date(2026, 9, 30))
    else:
        print("[X] --donem verin (ör. --donem 01.07.2026 30.09.2026)")
        return 2
    try:
        s = calistir(a.hareketler, a.cikti, donem, cari, vade=a.vade, firma=a.firma, sure=a.cevap_suresi, ayri=a.ayri_dosya)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    b = sum((c.bakiye for c in s["cariler"] if c.bakiye > 0), SIFIR)
    al = sum((-c.bakiye for c in s["cariler"] if c.bakiye < 0), SIFIR)
    print(f"[OK] {len(s['cariler'])} cari · borç bakiyeleri {tl(b)} TL · alacak bakiyeleri {tl(al)} TL")
    for u in s["uyarilar"]:
        if u["onem"] != "Bilgi":
            print(f"[!] {u['cari']} {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}" + (f" · {len(s['dosyalar'])} ayrı ekstre: {s['dosyalar'][0].parent.resolve()}" if s["dosyalar"] else ""))
    return 0


if __name__ == "__main__":
    sys.exit(main())
