import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
DONEM = (date(2026, 7, 1), date(2026, 9, 30))


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "c.xlsx"
        cls.s = main.calistir(ORNEK / "hareketler.csv", cls.cikti, DONEM, ORNEK / "cariler.csv", firma="Örnek AŞ", ayri=True)
        cls.c = {c.kod: c for c in cls.s["cariler"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_bakiyeler(self):
        c = self.c["120.001"]
        self.assertEqual((c.devir, c.borc, c.alacak, c.bakiye), (D("84000.00"), D("368450.00"), D("317390.00"), D("135060.00")))
        self.assertEqual(self.c["320.001"].bakiye, D("-375050.00"))
        self.assertEqual(self.c["120.003"].bakiye, D("-6000.00"))

    def test_fifo_ve_vade(self):
        acik = [(h.no, t) for h, t in self.c["120.002"].acik]
        self.assertEqual(acik, [("ABC2026000000390", D("95000.00")), ("ABC2026000000462", D("188000.00")), ("ABC2026000000536", D("96400.00"))])
        self.assertEqual([(h.no, t) for h, t in self.c["120.001"].acik], [("ABC2026000000531", D("135060.00"))])
        self.assertIsNone(next(h for h in self.c["120.001"].hareketler if h.tur == "Tahsilat").vade)   # ödemeye vade yazılmaz
        self.assertEqual(self.c["120.004"].hareketler[0].vade, date(2026, 4, 2))                     # cari vade 30 gün

    def test_uyarilar(self):
        u = {(x["tur"], x["cari"]) for x in self.s["uyarilar"]}
        self.assertTrue({("Mükerrer kayıt", "320.002"), ("Hareketsiz bakiye", "120.004"), ("Ters bakiye", "120.003"), ("Vadesi geçmiş", "120.002")} <= u)
        self.assertNotIn(("Vadesi geçmiş", "120.003"), u)                                           # fazla tahsilat vadesi geçmiş sayılmaz
        g = next(x for x in self.s["uyarilar"] if x["tur"] == "Vadesi geçmiş" and x["cari"] == "120.002")
        self.assertIn("283.000,00 TL", g["aciklama"])

    def test_metin_ve_dosyalar(self):
        m = main.mutabakat_metni(self.c["320.001"], DONEM[1], "Örnek AŞ", 15)
        self.assertIn("375.050,00 TL şirketimizden alacak bakiyesi", m)
        self.assertEqual(len(self.s["dosyalar"]), 6)
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Bakiye Özeti", "Ekstreler", "Açık Kalemler", "Mutabakat Metinleri", "Uyarılar"])
        self.assertEqual(wb["Ekstreler"]["H11"].value, 135060)


class KuralTesti(unittest.TestCase):
    def test_kartsiz_ve_varsayilan_vade(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t / "h.csv").write_text("Tarih;Cari Kod;Cari Unvan;Belge Türü;Belge No;Borç;Alacak\n01.06.2026;X1;X Ltd;Fatura;F1;1000;\n"
                                     "01.08.2026;X1;X Ltd;Tahsilat;T1;;400\n", encoding="utf-8")
            s = main.calistir(t / "h.csv", t / "o.xlsx", DONEM, vade=60)
            c = s["cariler"][0]
            self.assertEqual((c.unvan, c.devir, c.bakiye), ("X Ltd", D("1000"), D("600")))
            self.assertEqual(c.acik[0][0].vade, date(2026, 7, 31))
            self.assertTrue(any(u["tur"] == "Vadesi geçmiş" for u in s["uyarilar"]))

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--hareketler", str(ORNEK / "hareketler.csv"), "--donem", "30.09.2026", "01.07.2026"]), 2)
            self.assertEqual(main.main(["--hareketler", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
