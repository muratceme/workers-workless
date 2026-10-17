"""
Müşteri Hasar Dosyası Takibi — Gerisi Bende kod bloğu
Teknik ve Operasyon › Hasar Sorumlusu

Müşterilerin açık hasar dosyalarını sigorta şirketi, eksper ve bekleyen evrak bazında takip eder.
  Dosya yaşı     = rapor tarihi − ihbar tarihi (yoksa hasar tarihi)
  Hareketsiz gün = rapor tarihi − son işlem tarihi
  Öncelik (sırayla): Süre aşımı (yaş > --azami-gun, varsayılan 45) · Acil (hareketsiz > 2 × --hareketsiz) ·
                     Takip et (hareketsiz > --hareketsiz, varsayılan 7) · Evrak bekleniyor · Normal
  Açık dosya     = aşaması "kapandı", "ödendi", "ret" veya "iptal" ile başlamayan dosya
Varsayılan 45 gün bir iç takip eşiğidir; tazminatın ödenme süresi branşa, belgelerin tamamlanma tarihine ve mevzuata göre değişir.
Kod sigorta şirketine veya ekspere bağlanmaz; yalnız kendi takip listenizi sıralar.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"no": ("dosya no", "hasar no", "takip no"), "musteri": ("musteri", "sigortali", "musteri adi"), "telefon": ("telefon", "cep telefonu"), "brans": ("brans", "urun"),
         "sirket": ("sigorta sirketi", "sirket"), "sirket_no": ("sirket dosya no", "sirket hasar no", "hasar dosya no"), "eksper": ("eksper", "eksper adi"),
         "hasar": ("hasar tarihi",), "ihbar": ("ihbar tarihi", "bildirim tarihi"), "asama": ("asama", "durum", "dosya durumu"), "son": ("son islem tarihi", "son islem", "son hareket"),
         "tutar": ("tahmini tutar", "tahmini hasar", "talep tutari", "tutar"), "evrak": ("bekleyen evraklar", "bekleyen evrak", "eksik evrak", "eksik evraklar")}
KAPALI = ("kapandi", "kapali", "odendi", "ret", "red", "iptal")
SIRA = ["Süre aşımı", "Acil", "Takip et", "Evrak bekleniyor", "Normal"]
Z = Decimal(0)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("no", "asama")):
        no = metin(r["no"])
        if not no:
            continue
        kim = f"{no} (satır {r['_satir']})"
        if no in gorulen:
            uy.append(uyari("Yüksek", "Mükerrer dosya", kim, "Dosya numarası tekrar ediyor; ilk satır kullanıldı."))
            continue
        gorulen.add(no)
        d = {"no": no, "musteri": metin(r["musteri"]), "telefon": metin(r["telefon"]), "brans": metin(r["brans"]), "sirket": metin(r["sirket"]) or "Belirtilmemiş",
             "sirket_no": metin(r["sirket_no"]), "eksper": metin(r["eksper"]), "asama": metin(r["asama"]) or "Belirtilmemiş", "tutar": para(r["tutar"]),
             "evrak": [e.strip() for e in re.split(r"[|;\n]", metin(r["evrak"])) if e.strip()]}
        d["acik"] = not katla(d["asama"]).startswith(KAPALI)
        for a, ad in (("hasar", "Hasar tarihi"), ("ihbar", "İhbar tarihi"), ("son", "Son işlem tarihi")):
            d[a] = tarih(r[a])
            if metin(r[a]) and not d[a]:
                uy.append(uyari("Orta", "Tarih okunamadı", kim, f"{ad} '{metin(r[a])}'."))
        if d["tutar"] is None and metin(r["tutar"]):
            uy.append(uyari("Orta", "Tutar okunamadı", kim, f"'{metin(r['tutar'])}'. Tutar boş sayıldı."))
        if d["acik"] and not d["ihbar"] and not d["hasar"]:
            uy.append(uyari("Yüksek", "Tarih bilgisi yok", kim, "İhbar ve hasar tarihi boş; dosya yaşı hesaplanamadı."))
        elif d["acik"] and not d["ihbar"]:
            uy.append(uyari("Orta", "İhbar tarihi yok", kim, "Dosya yaşı hasar tarihinden hesaplandı. Şirkete ihbarın yapıldığını doğrulayın."))
        l.append(d)
    return l, uy


def aksiyon(d: dict, gecikmis: bool) -> str:
    a = []
    if d["evrak"]:
        a.append("Müşteriden evrak isteyin: " + ", ".join(d["evrak"]) + ".")
    if gecikmis:
        k = katla(d["asama"])
        if "eksper" in k:
            a.append(f"Eksperi arayın{' (' + d['eksper'] + ')' if d['eksper'] else ''}; rapor durumunu sorun.")
        elif "odeme" in k:
            a.append("Sigorta şirketinden ödeme tarihini sorun.")
        elif "evrak" not in k or not d["evrak"]:
            a.append("Sigorta şirketini arayın; dosyanın durumunu sorun.")
    if d["acik"] and not d["eksper"] and "ihbar" not in katla(d["asama"]) and not a:
        a.append("Eksper bilgisi girilmemiş; atama yapıldıysa dosyaya ekleyin.")
    return " ".join(a)


def takip(dosyalar: list[dict], rapor: date, hareketsiz: int = 7, azami_gun: int = 45) -> dict:
    acik = [d for d in dosyalar if d["acik"]]
    for d in acik:
        bas = d["ihbar"] or d["hasar"]
        d["yas"] = (rapor - bas).days if bas else None
        d["hareketsiz"] = (rapor - d["son"]).days if d["son"] else d["yas"]
        h = d["hareketsiz"]
        d["oncelik"] = ("Süre aşımı" if d["yas"] is not None and d["yas"] > azami_gun else "Acil" if h is not None and h > 2 * hareketsiz else "Takip et" if h is not None and h > hareketsiz
                        else "Evrak bekleniyor" if d["evrak"] else "Normal")
        d["aksiyon"] = aksiyon(d, d["oncelik"] in ("Süre aşımı", "Acil", "Takip et"))
    acik.sort(key=lambda d: (SIRA.index(d["oncelik"]), -(d["hareketsiz"] or 0), d["no"]))

    def grupla(alan: str) -> list[dict]:
        g = {}
        for d in acik:
            x = g.setdefault(d[alan] or "Atanmamış", {"ad": d[alan] or "Atanmamış", "dosya": 0, "yaslar": [], "tutar": Z, "evrakli": 0} | {s: 0 for s in SIRA})
            x["dosya"] += 1
            x["tutar"] += d["tutar"] or Z
            x[d["oncelik"]] += 1
            x["evrakli"] += bool(d["evrak"])
            if d["yas"] is not None:
                x["yaslar"].append(d["yas"])
        for x in g.values():
            x["ort_yas"] = yuvarla(Decimal(sum(x["yaslar"])) / len(x["yaslar"])) if x["yaslar"] else None
            x["en_eski"] = max(x["yaslar"]) if x["yaslar"] else None
        return sorted(g.values(), key=lambda x: (-x["dosya"], x["ad"]))

    evrak = [{"d": d, "evrak": e} for d in acik for e in d["evrak"]]
    return {"acik": acik, "kapali": len(dosyalar) - len(acik), "sirket": grupla("sirket"), "eksper": grupla("eksper"), "evrak": evrak, "tutar": sum((d["tutar"] or Z for d in acik), Z),
            "sayac": {s: sum(d["oncelik"] == s for d in acik) for s in SIRA}, "esik": (hareketsiz, azami_gun)}


def evrak_metni(r: dict) -> str:
    m = {}
    for x in r["evrak"]:
        m.setdefault((x["d"]["musteri"], x["d"]["telefon"]), {}).setdefault(x["d"]["no"], (x["d"], []))[1].append(x["evrak"])
    if not m:
        return "Bekleyen evrak yok.\n"
    b = []
    for (ad, tel), dosyalar in m.items():
        b.append(f"=== {ad or 'Müşteri belirtilmemiş'}" + (f" · {tel}" if tel else "") + " ===")
        for no, (d, liste) in dosyalar.items():
            b.append(f"{d['brans'] or 'Hasar'} dosyanız ({no}" + (f", {d['sirket']}" if d["sirket"] != "Belirtilmemiş" else "") + ") için bekleyen evraklar:")
            b += [f"  - {e}" for e in liste]
        b.append("")
    return "\n".join(b)


def rapor_yaz(cikti: Path, r: dict, rapor: date, uy: list[dict]) -> None:
    wb = Workbook()
    renk = {"Süre aşımı": KIRMIZI, "Acil": KIRMIZI, "Takip et": SARI, "Evrak bekleniyor": MAVI, "Normal": None}
    ws = wb.active
    ws.title = "Açık Dosyalar"
    ws.append([f"Rapor tarihi {rapor:%d.%m.%Y} · {len(r['acik'])} açık, {r['kapali']} kapalı dosya · tahmini tutar {r['tutar']:,.2f} · eşikler: hareketsiz {r['esik'][0]} gün, azami {r['esik'][1]} gün"])
    ws["A1"].font = Font(bold=True)
    _baslik(ws, ["Öncelik", "Dosya No", "Müşteri", "Branş", "Sigorta Şirketi", "Şirket Dosya No", "Eksper", "Aşama", "İhbar Tarihi", "Dosya Yaşı (gün)", "Son İşlem", "Hareketsiz (gün)",
                 "Tahmini Tutar", "Bekleyen Evrak", "Önerilen Aksiyon", "Not (siz doldurun)"], (16, 12, 26, 12, 20, 14, 18, 22, 12, 9, 12, 10, 16, 34, 60, 30), "A3")
    for d in r["acik"]:
        _satir(ws, [d["oncelik"], d["no"], d["musteri"], d["brans"], d["sirket"], d["sirket_no"], d["eksper"], d["asama"], d["ihbar"] or d["hasar"], d["yas"], d["son"], d["hareketsiz"], d["tutar"],
                    ", ".join(d["evrak"]), d["aksiyon"], ""], para_s=(13,), tarih_s=(9, 11), dolgu={1: renk[d["oncelik"]]}, sar=(14, 15))
    for ad, veri, kol in (("Şirket Bazında", r["sirket"], "Sigorta Şirketi"), ("Eksper Bazında", r["eksper"], "Eksper")):
        ws = wb.create_sheet(ad)
        _baslik(ws, [kol, "Açık Dosya", "Ortalama Yaş (gün)", "En Eski (gün)", *SIRA, "Evrak Bekleyen", "Tahmini Tutar"], (26, 9, 11, 10, 11, 8, 10, 12, 9, 10, 16))
        for x in veri:
            _satir(ws, [x["ad"], x["dosya"], x["ort_yas"], x["en_eski"], *[x[s] or None for s in SIRA], x["evrakli"] or None, x["tutar"]], para_s=(3, 11),
                   dolgu={5: KIRMIZI if x["Süre aşımı"] else None, 6: KIRMIZI if x["Acil"] else None})
    ws = wb.create_sheet("Bekleyen Evraklar")
    _baslik(ws, ["Dosya No", "Müşteri", "Telefon", "Sigorta Şirketi", "Evrak", "Dosya Yaşı (gün)", "İstendi (tarih)", "Geldi (tarih)"], (12, 26, 16, 20, 36, 9, 14, 14))
    for x in r["evrak"]:
        _satir(ws, [x["d"]["no"], x["d"]["musteri"], x["d"]["telefon"], x["d"]["sirket"], x["evrak"], x["d"]["yas"], "", ""])
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)
    (cikti.parent / "evrak_hatirlatma.txt").write_text(evrak_metni(r), encoding="utf-8")


def calistir(yol: Path, cikti: Path, rapor: date, hareketsiz: int = 7, azami_gun: int = 45) -> dict:
    dosyalar, uy = oku(yol)
    if not dosyalar:
        raise ValueError(f"{yol.name}: takip edilecek dosya yok.")
    r = takip(dosyalar, rapor, hareketsiz, azami_gun)
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, rapor, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Açık hasar dosyalarını sigorta şirketi, eksper ve bekleyen evrak bazında takip eder.")
    p.add_argument("--dosyalar", type=Path, default=ORNEK / "hasar_dosyalari.csv",
                   help="Dosya No, Müşteri, Telefon, Branş, Sigorta Şirketi, Şirket Dosya No, Eksper, Hasar Tarihi, İhbar Tarihi, Aşama, Son İşlem Tarihi, Tahmini Tutar, Bekleyen Evraklar")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 01.10.2026)")
    p.add_argument("--hareketsiz", type=int, default=7, help="Bu günden uzun süredir işlem görmeyen dosyalar takibe alınır (varsayılan 7)")
    p.add_argument("--azami-gun", type=int, default=45, help="Bu günden eski açık dosyalar 'Süre aşımı' olarak işaretlenir (varsayılan 45)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "hasar_dosyasi_takibi.xlsx")
    a = p.parse_args(argv)
    if not a.dosyalar.exists():
        print(f"[X] Dosya bulunamadı: {a.dosyalar}")
        return 1
    rapor = tarih(a.tarih) if a.tarih else (date(2026, 10, 1) if a.dosyalar == ORNEK / "hasar_dosyalari.csv" else date.today())
    if not rapor or a.hareketsiz < 1 or a.azami_gun < 1:
        print("[X] --tarih GG.AA.YYYY biçiminde; --hareketsiz ve --azami-gun en az 1 olmalı")
        return 2
    try:
        r = calistir(a.dosyalar, a.cikti, rapor, a.hareketsiz, a.azami_gun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(r['acik'])} açık dosya · " + " · ".join(f"{s}: {n}" for s, n in r["sayac"].items() if n) + f" · {len(r['evrak'])} bekleyen evrak")
    for d in r["acik"]:
        if d["oncelik"] in ("Süre aşımı", "Acil"):
            print(f"[!] {d['no']} · {d['oncelik']} · {d['sirket']}: yaş {d['yas']} gün, hareketsiz {d['hareketsiz']} gün")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[OK] Evrak hatırlatma metni: {(a.cikti.parent / 'evrak_hatirlatma.txt').resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
