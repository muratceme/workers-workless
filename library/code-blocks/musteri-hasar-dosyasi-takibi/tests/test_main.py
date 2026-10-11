import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
RAPOR = date(2026, 10, 1)


def dosya(no="D1", yas=10, hareketsiz=3, asama="Şirket incelemesinde", evrak=(), sirket="X Sigorta", eksper="Eksper A", tutar=1000, acik=True, musteri="M", ihbar=True):
    return {"no": no, "musteri": musteri, "telefon": "0500", "brans": "Kasko", "sirket": sirket, "sirket_no": "", "eksper": eksper, "asama": asama, "tutar": None if tutar is None else D(tutar),
            "evrak": list(evrak), "acik": acik, "hasar": RAPOR - timedelta(days=yas + 1), "ihbar": RAPOR - timedelta(days=yas) if ihbar else None,
            "son": None if hareketsiz is None else RAPOR - timedelta(days=hareketsiz)}


def tek(**kw):
    return main.takip([dosya(**kw)], RAPOR)["acik"][0]


class TakipTesti(unittest.TestCase):
    def test_oncelik_sinirlari(self):
        o = lambda **kw: tek(**kw)["oncelik"]   # noqa: E731
        self.assertEqual([o(yas=y) for y in (45, 46)], ["Normal", "Süre aşımı"])
        self.assertEqual([o(yas=20, hareketsiz=h) for h in (7, 8, 14, 15)], ["Normal", "Takip et", "Takip et", "Acil"])
        self.assertEqual(o(evrak=["Ruhsat"]), "Evrak bekleniyor")
        self.assertEqual(o(evrak=["Ruhsat"], hareketsiz=9), "Takip et")
        self.assertEqual(o(yas=60, hareketsiz=30, evrak=["Ruhsat"]), "Süre aşımı")
        d = tek(yas=12, hareketsiz=None)                                                    # son işlem yoksa dosya yaşı kullanılır
        self.assertEqual((d["hareketsiz"], d["oncelik"]), (12, "Takip et"))
        d = tek(yas=5, ihbar=False)                                                         # ihbar yoksa hasar tarihinden
        self.assertEqual(d["yas"], 6)
        self.assertEqual(main.takip([dosya(hareketsiz=5)], RAPOR, hareketsiz=3, azami_gun=5)["acik"][0]["oncelik"], "Süre aşımı")

    def test_aksiyon_metni(self):
        self.assertEqual(tek(evrak=["Ruhsat", "Tutanak"], asama="Evrak bekleniyor")["aksiyon"], "Müşteriden evrak isteyin: Ruhsat, Tutanak.")
        self.assertEqual(tek(hareketsiz=9, asama="Eksper incelemesinde")["aksiyon"], "Eksperi arayın (Eksper A); rapor durumunu sorun.")
        self.assertEqual(tek(hareketsiz=9, asama="Ödeme bekleniyor")["aksiyon"], "Sigorta şirketinden ödeme tarihini sorun.")
        self.assertEqual(tek(hareketsiz=9)["aksiyon"], "Sigorta şirketini arayın; dosyanın durumunu sorun.")
        self.assertEqual(tek(hareketsiz=9, asama="Evrak bekleniyor", evrak=["Ruhsat"])["aksiyon"], "Müşteriden evrak isteyin: Ruhsat.")
        self.assertEqual(tek(hareketsiz=9, asama="Evrak bekleniyor")["aksiyon"], "Sigorta şirketini arayın; dosyanın durumunu sorun.")
        self.assertEqual(tek(eksper="")["aksiyon"], "Eksper bilgisi girilmemiş; atama yapıldıysa dosyaya ekleyin.")
        self.assertEqual((tek(eksper="", asama="İhbar yapıldı")["aksiyon"], tek()["aksiyon"]), ("", ""))

    def test_gruplar_ve_evrak(self):
        l = [dosya("D1", 10, 3, sirket="A"), dosya("D2", 60, 2, sirket="A", evrak=["X", "Y"]), dosya("D3", 20, 20, sirket="B", eksper="", tutar=None), dosya("D4", acik=False, sirket="A")]
        r = main.takip(l, RAPOR)
        self.assertEqual(([d["no"] for d in r["acik"]], r["kapali"], r["tutar"]), (["D2", "D3", "D1"], 1, D(2000)))
        self.assertEqual([(x["ad"], x["dosya"], x["ort_yas"], x["en_eski"], x["Süre aşımı"], x["Acil"], x["evrakli"]) for x in r["sirket"]],
                         [("A", 2, D("35.00"), 60, 1, 0, 1), ("B", 1, D("20.00"), 20, 0, 1, 0)])
        self.assertEqual([(x["ad"], x["dosya"]) for x in r["eksper"]], [("Eksper A", 2), ("Atanmamış", 1)])
        self.assertEqual(([(x["d"]["no"], x["evrak"]) for x in r["evrak"]], r["sayac"]["Normal"]), ([("D2", "X"), ("D2", "Y")], 1))
        m = main.evrak_metni(r)
        self.assertIn("=== M · 0500 ===\nKasko dosyanız (D2, A) için bekleyen evraklar:\n  - X\n  - Y", m)
        self.assertEqual(main.evrak_metni(main.takip([dosya()], RAPOR)), "Bekleyen evrak yok.\n")


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "h.xlsx"
            r = main.calistir(ORNEK / "hasar_dosyalari.csv", cikti, RAPOR)
            self.assertEqual({d["no"]: (d["oncelik"], d["yas"], d["hareketsiz"]) for d in r["acik"]},
                             {"D-03": ("Süre aşımı", 60, 6), "D-05": ("Süre aşımı", 50, 11), "D-02": ("Acil", 29, 21), "D-07": ("Takip et", 12, 12), "D-01": ("Evrak bekleniyor", 10, 7),
                              "D-08": ("Evrak bekleniyor", 3, 2), "D-04": ("Normal", 5, 5)})
            self.assertEqual(([d["no"] for d in r["acik"]][:3], r["kapali"], r["tutar"], len(r["evrak"])), (["D-05", "D-03", "D-02"], 1, D(1192000), 5))
            self.assertEqual({x["ad"].split()[0]: (x["dosya"], x["ort_yas"]) for x in r["sirket"]}, {"Hayali": (3, D("12.33")), "Kurgu": (2, D("35.00")), "Uydurma": (2, D("31.00"))})
            self.assertIn("ödeme tarihini sorun", next(d["aksiyon"] for d in r["acik"] if d["no"] == "D-05"))
            self.assertEqual([u["tur"] for u in r["uyarilar"]], ["İhbar tarihi yok"])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Açık Dosyalar", "Şirket Bazında", "Eksper Bazında", "Bekleyen Evraklar", "Uyarılar"])
            self.assertIn("Sağlık dosyanız (D-08, Hayali Sigorta)", (Path(tmp) / "evrak_hatirlatma.txt").read_text(encoding="utf-8"))

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "d.csv"
            y.write_text("Dosya No;Aşama;Hasar Tarihi;İhbar Tarihi;Tahmini Tutar;Bekleyen Evraklar\nD1;ÖDENDİ;;;;\nD1;Açık;;;;\nD2;Reddedildi;;;;\nD3;İnceleme;dün;;çok;A; B\nD4;;01.09.2026;02.09.2026;5;\n",
                         encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual([(d["no"], d["acik"], d["asama"]) for d in l], [("D1", False, "ÖDENDİ"), ("D2", False, "Reddedildi"), ("D3", True, "İnceleme"), ("D4", True, "Belirtilmemiş")])
            self.assertEqual(sorted(u["tur"] for u in uy), ["Mükerrer dosya", "Tarih bilgisi yok", "Tarih okunamadı", "Tutar okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--dosyalar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "01.10.2026", "--hareketsiz", "3", "--azami-gun", "20"]), 0)
            self.assertEqual(main.main(arg + ["--tarih", "dün"]), 2)
            self.assertEqual(main.main(arg + ["--hareketsiz", "0"]), 2)
            self.assertEqual(main.main(["--dosyalar", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Dosya No;Aşama\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
