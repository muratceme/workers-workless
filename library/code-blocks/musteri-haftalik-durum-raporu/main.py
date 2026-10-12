"""
Müşteri Haftalık Durum Raporu — Workers / Workless kod bloğu
Müşteri Temsilciliği (Merchandising) › Production Merchandiser

Siparişlerin üretim aşaması ve sevk tarihi durumunu müşteriye gönderilecek haftalık rapor biçimine çevirir.
  Durum (sırayla denenir):
    Sevk edildi  = aşama "Sevk edildi" veya sevk edilen adet ≥ sipariş adedi
    Kısmi sevk   = 0 < sevk edilen adet < sipariş adedi
    Gecikmeli    = güncel sevk tarihi > planlanan sevk tarihi   (gecikme = aradaki gün);
                   güncel tarih yoksa ve planlanan tarih geçtiyse gecikme = rapor tarihi − planlanan tarih
    Riskli       = planlanan sevke --risk-gun (7) gün veya daha az kaldı ve aşama --risk-asama (Dikim) veya öncesinde
    Zamanında    = diğerleri
  Aşama sırası: Kumaş bekleniyor → Kesim → Dikim → Yıkama → Ütü-paket → Final kontrol → Sevk edildi
Rapor metni Türkçe veya İngilizce (--dil en) üretilir. Kod e-posta göndermez; metni siz gözden geçirip gönderirsiniz.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"musteri": ("musteri", "musteri adi", "buyer"), "po": ("siparis no", "po", "po no", "order no"), "model": ("model", "stil", "style"), "renk": ("renk", "color"),
         "adet": ("siparis adedi", "adet", "qty"), "asama": ("asama", "uretim asamasi", "durum"), "plan": ("planlanan sevk", "planlanan sevk tarihi", "termin", "ex factory"),
         "guncel": ("tahmini sevk", "guncel sevk", "guncel sevk tarihi", "tahmini sevk tarihi", "revize sevk"), "sevk": ("sevk edilen adet", "sevk adedi", "sevk edilen"), "not": ("not", "aciklama")}
ASAMALAR = [("kumas", "Kumaş bekleniyor", "Fabric awaited"), ("kesim", "Kesim", "Cutting"), ("dikim", "Dikim", "Sewing"), ("yika", "Yıkama", "Washing"), ("utu", "Ütü-paket", "Ironing & packing"),
            ("final", "Final kontrol", "Final inspection"), ("sevk edil", "Sevk edildi", "Shipped")]
EK_ANAHTAR = {"paket": 4, "kontrol": 5, "sevke hazir": 5, "dikis": 2, "bant": 2, "yuklendi": 6, "gonderildi": 6}
DURUM_EN = {"Sevk edildi": "Shipped", "Kısmi sevk": "Partially shipped", "Gecikmeli": "Delayed", "Riskli": "At risk", "Zamanında": "On track"}
SIRA = ["Gecikmeli", "Riskli", "Kısmi sevk", "Zamanında", "Sevk edildi"]
Z = Decimal(0)


def asama_sirasi(x) -> int | None:
    k = katla(x)
    for i, (anahtar, _, _) in enumerate(ASAMALAR):
        if anahtar in k:
            return i
    return next((i for a, i in EK_ANAHTAR.items() if a in k), None)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("po", "plan")):
        po = metin(r["po"])
        if not po:
            continue
        kim = f"{po} (satır {r['_satir']})"
        plan, guncel = tarih(r["plan"]), tarih(r["guncel"])
        if not plan:
            uy.append(uyari("Yüksek", "Planlanan sevk tarihi okunamadı", kim, f"'{metin(r['plan'])}'. Sipariş rapora alınmadı."))
            continue
        if metin(r["guncel"]) and not guncel:
            uy.append(uyari("Orta", "Güncel sevk tarihi okunamadı", kim, f"'{metin(r['guncel'])}'. Boş sayıldı."))
        a = asama_sirasi(r["asama"])
        if a is None:
            uy.append(uyari("Orta", "Aşama tanınmadı", kim, f"'{metin(r['asama'])}'. Aşama sırası bilinmediği için risk değerlendirmesi yapılmadı; rapora yazıldığı gibi alındı."))
        l.append({"musteri": metin(r["musteri"]) or "Müşteri", "po": po, "model": metin(r["model"]), "renk": metin(r["renk"]), "adet": para(r["adet"]), "asama": metin(r["asama"]) or "—",
                  "asama_no": a, "plan": plan, "guncel": guncel, "sevk": para(r["sevk"]) or Z, "not": metin(r["not"])})
    return l, uy


def degerlendir(siparisler: list[dict], rapor: date, risk_gun: int = 7, risk_asama: int = 2) -> dict:
    if not siparisler:
        raise ValueError("Raporlanacak sipariş yok.")
    for s in siparisler:
        s["gecikme"], s["kalan"] = None, (s["plan"] - rapor).days
        tam = s["adet"] is not None and s["adet"] > 0 and s["sevk"] >= s["adet"]
        if s["asama_no"] == 6 or tam:
            s["durum"] = "Sevk edildi"
        elif s["sevk"] > 0:
            s["durum"] = "Kısmi sevk"
        elif s["guncel"] and s["guncel"] > s["plan"]:
            s["durum"], s["gecikme"] = "Gecikmeli", (s["guncel"] - s["plan"]).days
        elif not s["guncel"] and s["plan"] < rapor:
            s["durum"], s["gecikme"] = "Gecikmeli", (rapor - s["plan"]).days
        elif s["kalan"] <= risk_gun and s["asama_no"] is not None and s["asama_no"] <= risk_asama:
            s["durum"] = "Riskli"
        else:
            s["durum"] = "Zamanında"
        s["tarih_yok"] = s["durum"] == "Gecikmeli" and not s["guncel"]
    siparisler.sort(key=lambda s: (s["musteri"], SIRA.index(s["durum"]), s["plan"], s["po"]))
    musteri = {}
    for s in siparisler:
        m = musteri.setdefault(s["musteri"], {"musteri": s["musteri"], "siparisler": [], "adet": Z} | {d: 0 for d in SIRA})
        m["siparisler"].append(s)
        m[s["durum"]] += 1
        m["adet"] += s["adet"] or Z
    return {"siparisler": siparisler, "musteri": list(musteri.values()), "say": {d: sum(s["durum"] == d for s in siparisler) for d in SIRA}}


def _sayi(x) -> str:
    return "—" if x is None else f"{int(x):,}".replace(",", ".")


def rapor_metni(r: dict, rapor: date, dil: str = "tr", imza: str = "") -> str:
    en, b = dil == "en", []
    for m in r["musteri"]:
        n = len(m["siparisler"])
        if en:
            b.append(f"=== {m['musteri']} · Weekly Status Report · {rapor:%d.%m.%Y} ===")
            b.append(f"{n} order(s): " + ", ".join(f"{m[d]} {DURUM_EN[d].lower()}" for d in SIRA if m[d]) + ".")
        else:
            b.append(f"=== {m['musteri']} · Haftalık Durum Raporu · {rapor:%d.%m.%Y} ===")
            b.append(f"Toplam {n} sipariş: " + ", ".join(f"{m[d]} {d.lower()}" for d in SIRA if m[d]) + ".")
        b.append("")
        for s in m["siparisler"]:
            asama = (ASAMALAR[s["asama_no"]][2] if en else ASAMALAR[s["asama_no"]][1]) if s["asama_no"] is not None else s["asama"]
            durum = DURUM_EN[s["durum"]] if en else s["durum"]
            if s["gecikme"]:
                durum += f" ({s['gecikme']} {'days' if en else 'gün'})"
            if s["durum"] == "Kısmi sevk":
                durum += f" ({_sayi(s['sevk'])} / {_sayi(s['adet'])})"
            b.append(" · ".join(x for x in (s["po"], s["model"], s["renk"], f"{_sayi(s['adet'])} {'pcs' if en else 'adet'}") if x))
            guncel = f"{s['guncel']:%d.%m.%Y}" if s["guncel"] else ("to be confirmed" if en else "bildirilecek") if s["tarih_yok"] else f"{s['plan']:%d.%m.%Y}"
            if en:
                b.append(f"  Stage: {asama} · Planned ex-factory: {s['plan']:%d.%m.%Y} · Current ex-factory: {guncel} · Status: {durum}")
            else:
                b.append(f"  Aşama: {asama} · Planlanan sevk: {s['plan']:%d.%m.%Y} · Güncel sevk: {guncel} · Durum: {durum}")
            if s["not"]:
                b.append(f"  {'Note' if en else 'Not'}: {s['not']}")
        b += ["", imza, ""] if imza else [""]
    return "\n".join(b)


def rapor_yaz(cikti: Path, r: dict, rapor: date, uy: list[dict], dil: str, imza: str) -> None:
    wb = Workbook()
    renk = {"Gecikmeli": KIRMIZI, "Riskli": SARI, "Kısmi sevk": MAVI, "Zamanında": None, "Sevk edildi": YESIL}
    ws = wb.active
    ws.title = "Sipariş Durumu"
    ws.append([f"Rapor tarihi {rapor:%d.%m.%Y} · {len(r['siparisler'])} sipariş · " + ", ".join(f"{n} {d.lower()}" for d, n in r["say"].items() if n)])
    ws["A1"].font = Font(bold=True)
    _baslik(ws, ["Müşteri", "Sipariş No", "Model", "Renk", "Sipariş Adedi", "Aşama", "Planlanan Sevk", "Güncel Sevk", "Kalan Gün", "Gecikme (gün)", "Sevk Edilen", "Durum", "Not",
                 "İç Not (müşteriye gitmez)"], (22, 14, 16, 12, 11, 20, 13, 13, 9, 10, 11, 14, 44, 30), "A3")
    for s in r["siparisler"]:
        _satir(ws, [s["musteri"], s["po"], s["model"], s["renk"], s["adet"], s["asama"], s["plan"], s["guncel"], s["kalan"] if s["durum"] != "Sevk edildi" else None, s["gecikme"], s["sevk"] or None,
                    s["durum"], s["not"] + (" Yeni sevk tarihi verilmemiş." if s["tarih_yok"] else ""), ""], tarih_s=(7, 8), dolgu={12: renk[s["durum"]]}, sar=(13,))
    ws = wb.create_sheet("Müşteri Özeti")
    _baslik(ws, ["Müşteri", "Sipariş", "Toplam Adet", *SIRA], (26, 9, 12, 11, 9, 11, 11, 12))
    for m in r["musteri"]:
        _satir(ws, [m["musteri"], len(m["siparisler"]), m["adet"], *[m[d] or None for d in SIRA]], dolgu={4: KIRMIZI if m["Gecikmeli"] else None, 5: SARI if m["Riskli"] else None})
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)
    (cikti.parent / "haftalik_durum_raporu.txt").write_text(rapor_metni(r, rapor, dil, imza), encoding="utf-8")


def calistir(yol: Path, cikti: Path, rapor: date, risk_gun: int = 7, risk_asama: int = 2, dil: str = "tr", imza: str = "") -> dict:
    siparisler, uy = oku(yol)
    r = degerlendir(siparisler, rapor, risk_gun, risk_asama)
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, rapor, uy, dil, imza)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Siparişlerin üretim aşaması ve sevk durumundan müşteriye gönderilecek haftalık raporu hazırlar.")
    p.add_argument("--siparisler", type=Path, default=ORNEK / "siparisler.csv",
                   help="Müşteri, Sipariş No, Model, Renk, Sipariş Adedi, Aşama, Planlanan Sevk, Tahmini Sevk, Sevk Edilen Adet, Not")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 12.10.2026)")
    p.add_argument("--risk-gun", type=int, default=7, help="Sevke bu kadar gün veya daha az kalan siparişler risk açısından değerlendirilir (varsayılan 7)")
    p.add_argument("--risk-asama", default="Dikim", help="Bu aşamada veya öncesindeki siparişler riskli sayılır (varsayılan Dikim)")
    p.add_argument("--dil", choices=("tr", "en"), default="tr", help="Rapor metninin dili (varsayılan tr)")
    p.add_argument("--imza", default="", help="Rapor metninin sonuna eklenecek imza")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "haftalik_durum_raporu.xlsx")
    a = p.parse_args(argv)
    if not a.siparisler.exists():
        print(f"[X] Dosya bulunamadı: {a.siparisler}")
        return 1
    rapor = tarih(a.tarih) if a.tarih else (date(2026, 10, 12) if a.siparisler == ORNEK / "siparisler.csv" else date.today())
    ra = asama_sirasi(a.risk_asama)
    if not rapor or a.risk_gun < 0 or ra is None:
        print("[X] --tarih GG.AA.YYYY biçiminde; --risk-gun sıfır veya pozitif; --risk-asama bilinen bir aşama olmalı (Kumaş, Kesim, Dikim, Yıkama, Ütü-paket, Final kontrol)")
        return 2
    try:
        r = calistir(a.siparisler, a.cikti, rapor, a.risk_gun, ra, a.dil, a.imza)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(r['siparisler'])} sipariş · {len(r['musteri'])} müşteri · " + ", ".join(f"{n} {d.lower()}" for d, n in r["say"].items() if n))
    for s in r["siparisler"]:
        if s["durum"] in ("Gecikmeli", "Riskli"):
            print(f"[!] {s['musteri']} · {s['po']} · {s['durum']}" + (f" ({s['gecikme']} gün)" if s["gecikme"] else "") + f" · aşama {s['asama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[OK] Rapor metni: {(a.cikti.parent / 'haftalik_durum_raporu.txt').resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
