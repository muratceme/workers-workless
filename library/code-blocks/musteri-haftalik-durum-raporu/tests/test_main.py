import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
RAPOR = date(2026, 10, 12)


def sp(po="PO1", asama="Dikim", plan=10, guncel=None, sevk=0, adet=1000, musteri="M", nt=""):
    t = lambda g: None if g is None else RAPOR + timedelta(days=g)   # noqa: E731
    return {"musteri": musteri, "po": po, "model": "TS", "renk": "", "adet": D(adet), "asama": asama, "asama_no": main.asama_sirasi(asama), "plan": t(plan), "guncel": t(guncel), "sevk": D(sevk),
            "not": nt}


def tek(**kw):
    return main.degerlendir([sp(**kw)], RAPOR)["siparisler"][0]


class DurumTesti(unittest.TestCase):
    def test_asama_tanima(self):
        self.assertEqual([main.asama_sirasi(x) for x in ("Kumaş bekleniyor", "KESİM", "Dikimde", "yıkama", "Ütü-paket", "Paketleme", "Final kontrol", "Sevk edildi", "Sevke hazır", "Nakış", "")],
                         [0, 1, 2, 3, 4, 4, 5, 6, 5, None, None])

    def test_durum_kurallari(self):
        self.assertEqual(tek(asama="Sevk edildi", plan=-7)["durum"], "Sevk edildi")
        self.assertEqual(tek(sevk=1000, plan=-7)["durum"], "Sevk edildi")
        self.assertEqual(tek(sevk=400)["durum"], "Kısmi sevk")
        x = tek(plan=3, guncel=10)
        self.assertEqual((x["durum"], x["gecikme"], x["tarih_yok"]), ("Gecikmeli", 7, False))
        x = tek(plan=-4)
        self.assertEqual((x["durum"], x["gecikme"], x["tarih_yok"]), ("Gecikmeli", 4, True))
        self.assertEqual(tek(plan=0)["durum"], "Riskli")                                         # sevk bugün: henüz gecikme değil, dikimde → riskli
        self.assertEqual([tek(plan=g)["durum"] for g in (7, 8)], ["Riskli", "Zamanında"])
        self.assertEqual([tek(plan=4, asama=a)["durum"] for a in ("Kesim", "Dikim", "Yıkama", "Final kontrol", "Nakış")], ["Riskli", "Riskli", "Zamanında", "Zamanında", "Zamanında"])
        self.assertEqual(tek(plan=4, guncel=4)["durum"], "Riskli")
        self.assertEqual(tek(plan=4, guncel=2, asama="Yıkama")["durum"], "Zamanında")
        self.assertEqual(main.degerlendir([sp(plan=10, asama="Yıkama")], RAPOR, risk_gun=14, risk_asama=3)["siparisler"][0]["durum"], "Riskli")

    def test_musteri_ozeti_ve_metin(self):
        l = [sp("PO1", plan=10), sp("PO2", plan=3, guncel=10, nt="Yıkama tekrarı"), sp("PO3", plan=-4), sp("PO4", "Ütü-paket", sevk=400), sp("PO5", "Sevk edildi", -5, musteri="N")]
        r = main.degerlendir(l, RAPOR)
        self.assertEqual([(m["musteri"], len(m["siparisler"]), m["Gecikmeli"], m["Zamanında"], m["Kısmi sevk"], m["adet"]) for m in r["musteri"]], [("M", 4, 2, 1, 1, D(4000)), ("N", 1, 0, 0, 0, D(1000))])
        self.assertEqual([s["po"] for s in r["siparisler"]], ["PO3", "PO2", "PO4", "PO1", "PO5"])
        tr = main.rapor_metni(r, RAPOR, "tr", "İmza")
        self.assertIn("=== M · Haftalık Durum Raporu · 12.10.2026 ===\nToplam 4 sipariş: 2 gecikmeli, 1 kısmi sevk, 1 zamanında.", tr)
        self.assertIn("PO2 · TS · 1.000 adet\n  Aşama: Dikim · Planlanan sevk: 15.10.2026 · Güncel sevk: 22.10.2026 · Durum: Gecikmeli (7 gün)\n  Not: Yıkama tekrarı", tr)
        self.assertIn("Güncel sevk: bildirilecek · Durum: Gecikmeli (4 gün)", tr)
        self.assertIn("Durum: Kısmi sevk (400 / 1.000)", tr)
        en = main.rapor_metni(r, RAPOR, "en")
        self.assertIn("=== M · Weekly Status Report · 12.10.2026 ===\n4 order(s): 2 delayed, 1 partially shipped, 1 on track.", en)
        self.assertIn("  Stage: Sewing · Planned ex-factory: 15.10.2026 · Current ex-factory: 22.10.2026 · Status: Delayed (7 days)", en)
        self.assertIn("Current ex-factory: to be confirmed", en)
        with self.assertRaises(ValueError):
            main.degerlendir([], RAPOR)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "h.xlsx"
            r = main.calistir(ORNEK / "siparisler.csv", cikti, RAPOR)
            self.assertEqual({s["po"]: (s["durum"], s["gecikme"]) for s in r["siparisler"]},
                             {"PO-501": ("Zamanında", None), "PO-502": ("Riskli", None), "PO-503": ("Gecikmeli", 7), "PO-504": ("Sevk edildi", None), "PO-601": ("Zamanında", None),
                              "PO-602": ("Gecikmeli", 4), "PO-603": ("Kısmi sevk", None), "PO-604": ("Zamanında", None)})
            self.assertEqual((r["say"], [u["tur"] for u in r["uyarilar"]]), ({"Gecikmeli": 2, "Riskli": 1, "Kısmi sevk": 1, "Zamanında": 3, "Sevk edildi": 1}, ["Aşama tanınmadı"]))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Sipariş Durumu", "Müşteri Özeti", "Uyarılar"])
            self.assertIn("PO-604 · KB-5006 · Bej · 400 adet\n  Aşama: Nakış", (Path(tmp) / "haftalik_durum_raporu.txt").read_text(encoding="utf-8"))

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "s.csv"
            y.write_text("Sipariş No;Aşama;Planlanan Sevk;Tahmini Sevk;Sipariş Adedi\nPO1;Dikim;20.10.2026;sonra;1.000\nPO2;Kesim;yakında;;\n;Kesim;20.10.2026;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual([(s["po"], s["plan"], s["guncel"], s["adet"], s["musteri"]) for s in l], [("PO1", date(2026, 10, 20), None, D(1000), "Müşteri")])
            self.assertEqual(sorted(u["tur"] for u in uy), ["Güncel sevk tarihi okunamadı", "Planlanan sevk tarihi okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--siparisler", str(y), "--cikti", c, "--tarih", "12.10.2026"]
            self.assertEqual(main.main(arg + ["--dil", "en", "--imza", "Best regards", "--risk-asama", "Yıkama", "--risk-gun", "10"]), 0)
            self.assertIn("Weekly Status Report", (Path(tmp) / "haftalik_durum_raporu.txt").read_text(encoding="utf-8"))
            self.assertEqual(main.main(arg + ["--risk-asama", "Boyama"]), 2)
            self.assertEqual(main.main(arg[:-1] + ["dün"]), 2)
            self.assertEqual(main.main(["--siparisler", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Sipariş No;Planlanan Sevk\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
