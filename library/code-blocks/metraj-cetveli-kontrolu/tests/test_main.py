import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def st(miktar, adet=None, boy=None, en=None, yuk=None, benzer=None, minha=False, poz="P1", tanim="T", satir=2, hata=False):
    f = lambda v: None if v is None else D(v)   # noqa: E731
    return {"poz": poz, "tanim": tanim, "satir": satir, "miktar": f(miktar), "minha": minha, "hata": hata, "benzer": f(benzer), "adet": f(adet), "boy": f(boy), "en": f(en), "yuk": f(yuk)}


def tek(*a, **kw):
    t = kw.pop("tolerans", D("0.01"))
    return main.denetle([st(*a, **kw)], tolerans=t)["satirlar"][0]


class DenetimTesti(unittest.TestCase):
    def test_sayi_cozme(self):
        self.assertEqual([main.sayi(x) for x in ("3,25", "3.25", "12.500", "1.250,50", "-7,20", 4, "üç", "")], [D("3.25"), D("3.25"), D("12.500"), D("1250.50"), D("-7.20"), D(4), None, None])

    def test_satir_hesabi(self):
        s = tek(6, 4, 2, "1.5", "0.5")
        self.assertEqual((s["hesap"], s["fark"], s["durum"]), (D("6.000"), D("0.000"), "Doğru"))
        s = tek("2.5", 3, 4, "0.25", "0.5", benzer=2)
        self.assertEqual((s["hesap"], s["fark"], s["durum"]), (D("3.000"), D("0.500"), "Hesap hatası"))
        self.assertEqual(tek(60, 1, 20, None, 3)["hesap"], D("60.000"))                                           # boş çarpan 1 sayılır
        self.assertEqual([(s["hesap"], s["minha"], s["durum"]) for s in (tek("-7.2", 4, "1.2", None, "1.5", minha=True), tek("-7.2", 4, "1.2", None, "1.5"), tek("7.2", 4, "1.2", None, "1.5", minha=True))],
                         [(D("-7.200"), True, "Doğru"), (D("-7.200"), True, "Doğru"), (D("-7.200"), True, "Hesap hatası")])
        self.assertEqual([tek(m, 1, "3.333", "3")["durum"] for m in ("10.00", "9.98")], ["Doğru", "Hesap hatası"])       # 9,999: fark 0,001 / 0,019
        self.assertEqual(tek("9.98", 1, "3.333", "3", tolerans=D("0.05"))["durum"], "Doğru")
        self.assertEqual((tek(None, 1, "12.5", None, 3)["durum"], tek(None, 1, "12.5", None, 3)["hesap"]), ("Miktar yazılmamış", D("37.500")))
        self.assertEqual((tek(5)["durum"], tek(5)["hesap"], tek(5, 2, hata=True)["durum"]), ("Doğru", D(5), "Hesaplanamadı"))

    def test_poz_toplami_mukerrer_ve_icmal(self):
        l = [st(6, 4, 2, "1.5", "0.5", satir=2), st("3.84", 8, "0.4", "0.4", 3, tanim="K1", satir=3), st("2.5", 3, 4, "0.25", "0.5", 2, satir=4), st("3.84", 8, "0.4", "0.4", 3, tanim="k1", satir=5),
             st(60, 1, 20, None, 3, poz="P2", satir=6), st("-7.2", 4, "1.2", None, "1.5", minha=True, poz="P2", satir=7), st(None, 1, "12.5", None, 3, poz="P2", satir=8)]
        r = main.denetle(l, {"p1": D("16.18"), "p2": D("90.30"), "p9": D(1)})
        self.assertEqual([(p["poz"], p["satir"], p["yazilan"], p["hesap"], p["fark"], p["hatali"], p["mukerrer"], p["eksik"], p["icmal_durum"]) for p in r["pozlar"]],
                         [("P1", 4, D("16.18"), D("16.680"), D("0.500"), 1, 2, 0, "İcmal farklı"), ("P2", 3, D("52.8"), D("90.300"), D("37.500"), 0, 0, 1, "Uyumlu")])
        self.assertEqual(([s["satir"] for s in r["satirlar"]], r["say"], r["cetvelde_yok"], r["icmalde_yok"]),
                         ([4, 8, 3, 5, 2, 6, 7], {"satir": 7, "hatali": 1, "mukerrer": 2, "hesaplanamayan": 0, "icmal_farkli": 1}, ["p9"], []))
        r = main.denetle([st(1, 1, tanim="A"), st(1, 1, tanim="A", poz="P2"), st(1, 1, 1, tanim="A")], {"p1": D(2)})
        self.assertEqual((r["say"]["mukerrer"], r["icmalde_yok"]), (0, ["P2"]))                                   # poz veya çarpan farklıysa mükerrer değil
        with self.assertRaises(ValueError):
            main.denetle([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "m.xlsx"
            r = main.calistir(ORNEK / "metraj_cetveli.csv", cikti, ORNEK / "icmal.csv")
            self.assertEqual(([(p["poz"], p["hesap"], p["icmal_durum"]) for p in r["pozlar"]], r["say"], r["uyarilar"]),
                             ([("POZ-01", D("16.680"), "İcmal farklı"), ("POZ-02", D("90.300"), "Uyumlu")], {"satir": 7, "hatali": 1, "mukerrer": 2, "hesaplanamayan": 0, "icmal_farkli": 1}, []))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Satır Kontrolü", "Poz Toplamları", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, ic = Path(tmp) / "m.csv", Path(tmp) / "i.csv"
            y.write_text("Poz No;Açıklama;Adet;Boy;En;Yükseklik;Miktar\nP1;A;2;3.50;;;7.00\nP1;B;iki;3;;;6\nP1;C;;;;;12,5\nP1;D;;;;;\n;E;1;1;;;1\nP1;F;1;2;;;çok\n", encoding="utf-8")
            ic.write_text("Poz;Toplam\nP1;10\nP1;9,5\nP2;yok\n", encoding="utf-8")
            l, icmal, uy = main.oku(y, ic)
            self.assertEqual(([(s["tanim"], s["adet"], s["boy"], s["miktar"], s["hata"]) for s in l], icmal),
                             ([("A", D(2), D("3.50"), D("7.00"), False), ("B", None, D(3), D(6), True), ("C", None, None, D("12.5"), False), ("F", D(1), D(2), None, False)], {"p1": D("19.5")}))
            self.assertEqual(sorted(u["tur"] for u in uy), ["Miktar okunamadı", "Çarpan okunamadı", "Çarpanı olmayan satır", "İcmal miktarı okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--metraj", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--icmal", str(ic), "--tolerans", "0,005"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(arg + ["--tolerans", "-1"]), 2)
            self.assertEqual(main.main(arg + ["--icmal", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Poz No;Miktar\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
