"""
Metraj Cetveli Kontrolü — Gerisi Bende kod bloğu
Teknik Ofis › Teknik Ofis Mühendisi

Metraj cetvelindeki hesap satırlarını (boyut × adet) yeniden hesaplar, poz toplamlarını ve mükerrer satırları kontrol eder.
  Hesaplanan miktar = benzer × adet × boy × en × yükseklik          (boş bırakılan çarpan 1 sayılır; "Minha" işaretli satır eksi alınır)
  Hesap hatası      = |hesaplanan − cetvelde yazılan miktar| > --tolerans (varsayılan 0,01)
  Poz toplamı       = pozun satırlarının hesaplanan miktarları toplamı; cetvelde yazılan toplamla ve verilirse icmal dosyasındaki miktarla karşılaştırılır
  Mükerrer satır    = aynı poz, aynı tanım ve aynı çarpanlarla birden fazla satır
Sayılarda virgül de nokta da ondalık ayırıcıdır (3,25 = 3.25); binlik ayırıcı yalnız "1.250,50" yazımında tanınır.
Kod aritmetiği denetler: ölçülerin projeye uygunluğunu, pozun doğru seçildiğini ve metraj yöntemini (boşluk düşme kuralları vb.) değerlendirmez.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"poz": ("poz no", "poz", "is kalemi no", "poz numarasi"), "tanim": ("tanim", "aciklama", "yer", "yapilan isin cinsi", "mahal"), "benzer": ("benzer", "benzer adedi", "tekrar"),
         "adet": ("adet", "sayi"), "boy": ("boy", "uzunluk"), "en": ("en", "genislik"), "yuk": ("yukseklik", "derinlik", "kalinlik"), "miktar": ("miktar", "tutar", "sonuc", "hesaplanan"),
         "minha": ("minha", "dusulecek", "cikarilacak")}
ICMAL = {"poz": SUTUN["poz"], "miktar": ("miktar", "toplam", "icmal miktari", "metraj")}
CARPAN = ("benzer", "adet", "boy", "en", "yuk")
EVET = {"e", "evet", "x", "1", "minha", "var"}
Z = Decimal(0)


def sayi(x) -> Decimal | None:
    """3,25 · 3.25 · 1.250,50 → Decimal; virgül de nokta da ondalık ayırıcıdır."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    s = s.replace(".", "").replace(",", ".") if "," in s and "." in s else s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def oku(yol: Path, icmal_yol: Path | None = None) -> tuple[list[dict], dict, list[dict]]:
    uy, l, icmal = [], [], {}
    for r in kayitlar(yol, SUTUN, ("poz",)):
        poz = metin(r["poz"])
        if not poz:
            continue
        kim = f"{poz} · {metin(r['tanim']) or 'satır'} (satır {r['_satir']})"
        c, hata = {}, False
        for a in CARPAN:
            c[a] = sayi(r[a])
            if c[a] is None and metin(r[a]):
                uy.append(uyari("Yüksek", "Çarpan okunamadı", kim, f"{SUTUN[a][0]} '{metin(r[a])}'. Satır yeniden hesaplanamadı."))
                hata = True
        m = sayi(r["miktar"])
        if m is None and metin(r["miktar"]):
            uy.append(uyari("Yüksek", "Miktar okunamadı", kim, f"'{metin(r['miktar'])}'. Yazılan miktar boş sayıldı."))
        if all(v is None for v in c.values()) and not hata:
            if m is None:
                continue
            uy.append(uyari("Orta", "Çarpanı olmayan satır", kim, f"Yalnız miktar ({m:f}) yazılmış; hesap satırı yok. Miktar olduğu gibi alındı, denetlenemedi."))
        l.append({"poz": poz, "tanim": metin(r["tanim"]), "satir": r["_satir"], "miktar": m, "minha": katla(r["minha"]) in EVET, "hata": hata} | c)
    if icmal_yol:
        for r in kayitlar(icmal_yol, ICMAL, ("poz", "miktar")):
            p, m = metin(r["poz"]), sayi(r["miktar"])
            if not p:
                continue
            if m is None:
                uy.append(uyari("Yüksek", "İcmal miktarı okunamadı", p, f"{icmal_yol.name} satır {r['_satir']}: '{metin(r['miktar'])}'."))
                continue
            icmal[katla(p)] = icmal.get(katla(p), Z) + m
    return l, icmal, uy


def denetle(satirlar: list[dict], icmal: dict | None = None, tolerans: Decimal = Decimal("0.01")) -> dict:
    if not satirlar:
        raise ValueError("Denetlenecek metraj satırı yok.")
    icmal, gorulen = icmal or {}, defaultdict(list)
    for s in satirlar:
        carpanlar = [s[a] for a in CARPAN if s[a] is not None]
        if s["hata"]:
            s["hesap"] = None
        elif not carpanlar:
            s["hesap"] = s["miktar"]
        else:
            h = Decimal(1)
            for c in carpanlar:
                h *= c
            if s["minha"] or (s["miktar"] is not None and s["miktar"] < 0 < h):          # eksi yazılmış miktar da minha sayılır
                h, s["minha"] = -abs(h), True
            s["hesap"] = yuvarla(h, "0.001")
        s["fark"] = s["hesap"] - s["miktar"] if s["hesap"] is not None and s["miktar"] is not None else None
        s["durum"] = "Hesaplanamadı" if s["hesap"] is None else "Miktar yazılmamış" if s["miktar"] is None else "Hesap hatası" if abs(s["fark"]) > tolerans else "Doğru"
        s["mukerrer"] = False
        gorulen[(katla(s["poz"]), katla(s["tanim"]), *[s[a] for a in CARPAN], s["minha"])].append(s)
    for grup in gorulen.values():
        if len(grup) > 1:
            for s in grup:
                s["mukerrer"] = True
    poz = {}
    for s in satirlar:
        p = poz.setdefault(katla(s["poz"]), {"poz": s["poz"], "satir": 0, "hesap": Z, "yazilan": Z, "hatali": 0, "mukerrer": 0, "eksik": 0})
        p["satir"] += 1
        p["hesap"] += s["hesap"] or Z
        p["yazilan"] += s["miktar"] or Z
        p["hatali"] += s["durum"] == "Hesap hatası"
        p["mukerrer"] += s["mukerrer"]
        p["eksik"] += s["durum"] in ("Hesaplanamadı", "Miktar yazılmamış")
    for k, p in poz.items():
        p["fark"] = p["hesap"] - p["yazilan"]
        p["icmal"] = icmal.get(k)
        p["icmal_fark"] = p["hesap"] - p["icmal"] if p["icmal"] is not None else None
        p["icmal_durum"] = "" if p["icmal"] is None else "Uyumlu" if abs(p["icmal_fark"]) <= tolerans else "İcmal farklı"
    sira = ("Hesap hatası", "Hesaplanamadı", "Miktar yazılmamış", "Doğru")
    return {"satirlar": sorted(satirlar, key=lambda s: (sira.index(s["durum"]), not s["mukerrer"], s["satir"])), "pozlar": list(poz.values()), "tolerans": tolerans,
            "icmalde_yok": [p["poz"] for k, p in poz.items() if icmal and k not in icmal], "cetvelde_yok": sorted(k for k in icmal if k not in poz),
            "say": {"satir": len(satirlar), "hatali": sum(s["durum"] == "Hesap hatası" for s in satirlar), "mukerrer": sum(s["mukerrer"] for s in satirlar),
                    "hesaplanamayan": sum(s["durum"] == "Hesaplanamadı" for s in satirlar), "icmal_farkli": sum(p["icmal_durum"] == "İcmal farklı" for p in poz.values())}}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    renk = {"Hesap hatası": KIRMIZI, "Hesaplanamadı": KIRMIZI, "Miktar yazılmamış": MAVI}
    ws = wb.active
    ws.title = "Satır Kontrolü"
    _baslik(ws, ["Durum", "Mükerrer", "Cetvel Satırı", "Poz No", "Tanım", "Minha", "Benzer", "Adet", "Boy", "En", "Yükseklik", "Yazılan Miktar", "Hesaplanan Miktar", "Fark", "Düzeltme (siz doldurun)"],
            (17, 9, 8, 16, 34, 7, 8, 8, 9, 9, 10, 13, 13, 11, 28))
    for s in r["satirlar"]:
        n = _satir(ws, [s["durum"], "Evet" if s["mukerrer"] else "", s["satir"], s["poz"], s["tanim"], "Evet" if s["minha"] else "", *[s[a] for a in CARPAN], s["miktar"], s["hesap"], s["fark"], ""],
                   dolgu={1: renk.get(s["durum"]), 2: SARI if s["mukerrer"] else None})
        for j in (12, 13, 14):
            ws.cell(n, j).number_format = "#,##0.000"
    ws.append([])
    ws.append([f"Hesaplanan = benzer × adet × boy × en × yükseklik (boş çarpan 1). Tolerans {r['tolerans']:f}. Kod aritmetiği denetler; ölçülerin projeye uygunluğunu ve metraj yöntemini değerlendirmez."])
    ws = wb.create_sheet("Poz Toplamları")
    _baslik(ws, ["Poz No", "Satır Sayısı", "Yazılan Toplam", "Hesaplanan Toplam", "Fark", "Hatalı Satır", "Mükerrer Satır", "Eksik Satır", "İcmal Miktarı", "İcmal Farkı", "İcmal Durumu"],
            (18, 8, 14, 14, 11, 9, 9, 9, 13, 11, 14))
    for p in r["pozlar"]:
        n = _satir(ws, [p["poz"], p["satir"], p["yazilan"], p["hesap"], p["fark"], p["hatali"] or None, p["mukerrer"] or None, p["eksik"] or None, p["icmal"], p["icmal_fark"], p["icmal_durum"]],
                   dolgu={5: KIRMIZI if abs(p["fark"]) > r["tolerans"] else None, 11: KIRMIZI if p["icmal_durum"] == "İcmal farklı" else YESIL if p["icmal_durum"] else None})
        for j in (3, 4, 5, 9, 10):
            ws.cell(n, j).number_format = "#,##0.000"
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, icmal_yol: Path | None = None, tolerans: Decimal = Decimal("0.01")) -> dict:
    satirlar, icmal, uy = oku(yol, icmal_yol)
    r = denetle(satirlar, icmal, tolerans)
    for p in r["icmalde_yok"]:
        uy.append(uyari("Orta", "Poz icmalde yok", p, "Metraj cetvelinde var, icmal dosyasında bu poz yok."))
    for p in r["cetvelde_yok"]:
        uy.append(uyari("Orta", "İcmaldeki pozun metrajı yok", p, "İcmal dosyasında var, metraj cetvelinde bu pozla satır yok."))
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Metraj cetvelindeki hesap satırlarını yeniden hesaplar; poz toplamlarını ve mükerrer satırları denetler.")
    p.add_argument("--metraj", type=Path, default=ORNEK / "metraj_cetveli.csv", help="Poz No, Tanım, Benzer, Adet, Boy, En, Yükseklik, Miktar, Minha")
    p.add_argument("--icmal", type=Path, help="İsteğe bağlı: Poz No, Miktar (icmaldeki toplamlar)")
    p.add_argument("--tolerans", default="0,01", help="Bu farka kadar satır doğru sayılır (varsayılan 0,01)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "metraj_kontrolu.xlsx")
    a = p.parse_args(argv)
    icmal, t = a.icmal or (ORNEK / "icmal.csv" if a.metraj == ORNEK / "metraj_cetveli.csv" else None), sayi(a.tolerans)
    if t is None or t < 0:
        print("[X] --tolerans sıfır veya pozitif sayı olmalı")
        return 2
    for y in (a.metraj, icmal):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.metraj, a.cikti, icmal, t)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    s = r["say"]
    print(f"[OK] {s['satir']} satır · {len(r['pozlar'])} poz · {s['hatali']} hesap hatası, {s['mukerrer']} mükerrer satır, {s['hesaplanamayan']} hesaplanamayan, {s['icmal_farkli']} icmali farklı poz")
    for x in r["satirlar"]:
        if x["durum"] == "Hesap hatası":
            print(f"[!] Satır {x['satir']} · {x['poz']} · {x['tanim']}: yazılan {x['miktar']:f}, hesaplanan {x['hesap']:f}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
