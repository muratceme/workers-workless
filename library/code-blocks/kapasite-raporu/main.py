"""
Kapasite Raporu — Workers / Workless kod bloğu
Bilgi Teknolojileri › Sistem Uzmanı

Disk, bellek ve işlemci kullanım ölçümlerinden büyüme eğilimini ve eşiğe ulaşma tarihini tahmin eder.
  Günlük büyüme       = ölçümlere en küçük kareler yöntemiyle uydurulan doğrunun eğimi (birim ÷ gün)
  Kullanım oranı      = son ölçüm ÷ son kapasite
  Eşiğe kalan gün     = (kapasite × --esik % − son ölçüm) ÷ günlük büyüme, yukarı yuvarlanır;  tarih = son ölçüm tarihi + gün
  Dolmaya kalan gün   = (kapasite − son ölçüm) ÷ günlük büyüme, yukarı yuvarlanır
  R²                  = doğrunun ölçümleri açıklama oranı (1 = tam doğrusal); 0,5 altı "eğilim zayıf" sayılır
  Durum               = Eşik aşıldı (oran ≥ eşik) · Kritik (eşiğe ≤ --kritik gün) · Uyarı (≤ --uyari gün) · Normal ·
                        Büyüme yok (eğim ≤ 0) · Yetersiz veri (ölçüm < --min-olcum veya tek tarih)
Tahmin geçmişin aynen süreceğini varsayar; toplu veri yükleme, temizlik veya yeni proje gibi kırılmaları bilemez.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
from fractions import Fraction

SUTUN = {"tarih": ("tarih", "olcum tarihi", "date"), "sistem": ("sistem", "sunucu", "server", "host"), "kaynak": ("kaynak", "disk", "birim adi", "resource", "volume"), "kullanilan": ("kullanilan", "kullanim", "dolu", "used"),
         "kapasite": ("kapasite", "toplam", "boyut", "capacity", "total"), "birim": ("birim", "olcu birimi", "unit")}
SIRA = {"Eşik aşıldı": 0, "Kritik": 1, "Uyarı": 2, "Normal": 3, "Büyüme yok": 4, "Yetersiz veri": 5}
D_RENK = {"Eşik aşıldı": KIRMIZI, "Kritik": KIRMIZI, "Uyarı": SARI, "Normal": YESIL}
YUZ = Decimal(100)


def olcu(x) -> Decimal | None:
    """Ölçü değerleri: virgül veya tek nokta ondalık ayırıcıdır (12.5 = 12,5; 1.250,5 = 1250,5)."""
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x or ""))
    s = s.replace(".", "").replace(",", ".") if "," in s else s
    try:
        return Decimal(s) if s else None
    except InvalidOperation:
        return None


def ondalik(f: Fraction, hane: str = "0.01") -> Decimal:
    return yuvarla(Decimal(f.numerator) / Decimal(f.denominator), hane)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("tarih", "sistem", "kullanilan", "kapasite")):
        s = metin(r["sistem"])
        if not s:
            continue
        t, k, kap = tarih(r["tarih"]), olcu(r["kullanilan"]), olcu(r["kapasite"])
        if t is None or k is None or k < 0 or kap is None or kap <= 0:
            uy.append(uyari("Yüksek", "Ölçüm okunamadı", f"{s} · {metin(r['kaynak'])} (satır {r['_satir']})", f"Tarih '{metin(r['tarih'])}', kullanılan '{metin(r['kullanilan'])}', kapasite '{metin(r['kapasite'])}'. Ölçüm rapora alınmadı."))
            continue
        l.append({"tarih": t, "sistem": s, "kaynak": metin(r["kaynak"]) or "(kaynak yok)", "kullanilan": k, "kapasite": kap, "birim": metin(r["birim"])})
    return l, uy


def egilim(noktalar: list[tuple[int, Decimal]]) -> tuple[Fraction, Fraction | None] | None:
    """(gün, değer) noktalarına doğru uydurur; (eğim, R²) döndürür. Tek tarih varsa None."""
    n = len(noktalar)
    xo, yo = Fraction(sum(x for x, _ in noktalar), n), sum((Fraction(y) for _, y in noktalar), Fraction(0)) / n
    sxx, sxy, syy = (sum(((a(x, y)) for x, y in noktalar), Fraction(0)) for a in (lambda x, y: (x - xo) ** 2, lambda x, y: (x - xo) * (Fraction(y) - yo), lambda x, y: (Fraction(y) - yo) ** 2))
    if sxx == 0:
        return None
    return sxy / sxx, (sxy * sxy / (sxx * syy) if syy else None)


def raporla(olcumler: list[dict], bugun: date, esik: Decimal = Decimal(90), kritik: int = 30, uyari_gun: int = 90, min_olcum: int = 3, bayat: int = 14) -> dict:
    if not olcumler:
        raise ValueError("Raporlanacak ölçüm yok.")
    uy, grup = [], defaultdict(dict)
    for o in sorted(olcumler, key=lambda o: o["tarih"]):
        grup[(katla(o["sistem"]), katla(o["kaynak"]))][o["tarih"]] = o      # aynı güne birden çok ölçüm varsa sonuncusu geçerlidir
    kaynaklar = []
    for gunler in grup.values():
        seri = [gunler[t] for t in sorted(gunler)]
        ilk, son = seri[0], seri[-1]
        kim = f"{son['sistem']} · {son['kaynak']}"
        k = {"sistem": son["sistem"], "kaynak": son["kaynak"], "birim": son["birim"], "olcum": len(seri), "ilk_tarih": ilk["tarih"], "son_tarih": son["tarih"], "son": son["kullanilan"], "kapasite": son["kapasite"],
             "oran": yuvarla(son["kullanilan"] * YUZ / son["kapasite"]), "bos": son["kapasite"] - son["kullanilan"], "gunluk": None, "aylik": None, "r2": None, "esik_gun": None, "esik_tarih": None, "dolma_gun": None, "dolma_tarih": None}
        e = egilim([((o["tarih"] - ilk["tarih"]).days, o["kullanilan"]) for o in seri]) if len(seri) >= min_olcum else None
        if len({o["kapasite"] for o in seri}) > 1:
            uy.append(uyari("Bilgi", "Kapasite değişmiş", kim, f"Ölçümler arasında kapasite değişmiş; hesapta son kapasite ({son['kapasite']} {son['birim']}) kullanıldı.".replace(" )", ")")))
        if (bugun - son["tarih"]).days > bayat:
            uy.append(uyari("Orta", "Ölçüm eski", kim, f"Son ölçüm {son['tarih']:%d.%m.%Y}; {(bugun - son['tarih']).days} gün önce. Tahmin güncel durumu yansıtmayabilir."))
        if e:
            egim, r2 = e
            k["gunluk"], k["aylik"], k["r2"] = ondalik(egim, "0.001"), ondalik(egim * 30), ondalik(r2) if r2 is not None else None
        esik_deger = son["kapasite"] * esik / YUZ
        if son["kullanilan"] >= esik_deger:
            k["durum"] = "Eşik aşıldı"
        elif e is None:
            k["durum"] = "Yetersiz veri"
        elif egim <= 0:
            k["durum"] = "Büyüme yok"
        if e and egim > 0:
            if son["kullanilan"] < esik_deger:
                k["esik_gun"] = math.ceil((Fraction(esik_deger) - Fraction(son["kullanilan"])) / egim)
                k["esik_tarih"] = son["tarih"] + timedelta(days=k["esik_gun"])
                k["durum"] = "Kritik" if k["esik_gun"] <= kritik else "Uyarı" if k["esik_gun"] <= uyari_gun else "Normal"
            if son["kullanilan"] < son["kapasite"]:
                k["dolma_gun"] = math.ceil((Fraction(son["kapasite"]) - Fraction(son["kullanilan"])) / egim)
                k["dolma_tarih"] = son["tarih"] + timedelta(days=k["dolma_gun"])
        if k["durum"] == "Eşik aşıldı":
            uy.append(uyari("Yüksek", "Eşik aşıldı", kim, f"Kullanım %{k['oran']} (eşik %{esik})." + (f" Bu hızla {k['dolma_tarih']:%d.%m.%Y} civarında ({k['dolma_gun']} gün) tamamen dolar." if k["dolma_tarih"] else "")))
        elif k["durum"] in ("Kritik", "Uyarı"):
            uy.append(uyari("Yüksek" if k["durum"] == "Kritik" else "Orta", f"Eşiğe {k['durum'] == 'Kritik' and kritik or uyari_gun} günden az kaldı", kim,
                            f"Kullanım %{k['oran']}; günlük büyüme {k['gunluk']} {k['birim']}. %{esik} eşiğine {k['esik_tarih']:%d.%m.%Y} civarında ({k['esik_gun']} gün) ulaşır.".replace("  ", " ")))
            if k["r2"] is not None and k["r2"] < Decimal("0.5"):
                uy.append(uyari("Bilgi", "Eğilim zayıf", kim, f"R² {k['r2']}; ölçümler doğrusal bir büyüme göstermiyor, tahmine temkinli yaklaşın."))
        elif k["durum"] == "Yetersiz veri":
            uy.append(uyari("Bilgi", "Yetersiz veri", kim, f"{k['olcum']} ölçüm var; eğilim için farklı tarihlerde en az {min_olcum} ölçüm gerekir."))
        kaynaklar.append(k)
    kaynaklar.sort(key=lambda k: (SIRA[k["durum"]], k["esik_gun"] if k["esik_gun"] is not None else 10 ** 9, katla(k["sistem"]), katla(k["kaynak"])))
    return {"kaynaklar": kaynaklar, "say": {d: sum(k["durum"] == d for k in kaynaklar) for d in SIRA}, "esik": esik, "bugun": bugun, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kapasite Özeti"
    _baslik(ws, ["Durum", "Sistem", "Kaynak", "Birim", "Kullanılan", "Kapasite", "Boş", "Kullanım %", "Günlük Büyüme", "30 Günlük Büyüme", "R²", f"%{r['esik']} Eşiğine Kalan Gün", "Eşik Tarihi", "Dolmaya Kalan Gün", "Dolma Tarihi",
                 "Ölçüm Sayısı", "İlk Ölçüm", "Son Ölçüm"], (13, 14, 18, 7, 13, 13, 13, 10, 11, 12, 7, 12, 12, 11, 12, 8, 12, 12))
    for k in r["kaynaklar"]:
        n = _satir(ws, [k["durum"], k["sistem"], k["kaynak"], k["birim"], k["son"], k["kapasite"], k["bos"], k["oran"], k["gunluk"], k["aylik"], k["r2"], k["esik_gun"], k["esik_tarih"], k["dolma_gun"], k["dolma_tarih"], k["olcum"],
                        k["ilk_tarih"], k["son_tarih"]], para_s=(5, 6, 7, 8, 10, 11), tarih_s=(13, 15, 17, 18), dolgu={1: D_RENK.get(k["durum"])})
        ws.cell(n, 9).number_format = "#,##0.000"
    ws.append([])
    ws.append([f"Durum tarihi {r['bugun']:%d.%m.%Y}. Tarihler son ölçümden itibaren doğrusal büyüme varsayımıyla hesaplanır; kesin değildir."])
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, esik: Decimal = Decimal(90), kritik: int = 30, uyari_gun: int = 90, min_olcum: int = 3, bayat: int = 14) -> dict:
    olcumler, uy = oku(yol)
    r = raporla(olcumler, bugun, esik, kritik, uyari_gun, min_olcum, bayat)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Kullanım ölçümlerinden büyüme eğilimi ve dolma tarihi tahmini üretir.")
    p.add_argument("--olcumler", type=Path, help="Tarih, Sistem, Kaynak, Kullanılan, Kapasite, Birim")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--esik", default="90", help="Kullanım eşiği, %% (varsayılan 90)")
    p.add_argument("--kritik", type=int, default=30, help="Eşiğe bu kadar gün veya daha az kaldıysa Kritik (varsayılan 30)")
    p.add_argument("--uyari", type=int, default=90, help="Eşiğe bu kadar gün veya daha az kaldıysa Uyarı (varsayılan 90)")
    p.add_argument("--min-olcum", type=int, default=3, help="Eğilim için gereken en az ölçüm (varsayılan 3)")
    p.add_argument("--bayat", type=int, default=14, help="Son ölçüm bu kadar günden eskiyse uyarılır (varsayılan 14)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kapasite_raporu.xlsx")
    a = p.parse_args(argv)
    ornek = a.olcumler is None
    yol = ORNEK / "olcumler.csv" if ornek else a.olcumler
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    esik = olcu(a.esik)
    if bugun is None or esik is None or not 0 < esik <= 100 or not 0 < a.kritik < a.uyari or a.min_olcum < 2 or a.bayat < 0:
        print("[X] --tarih GG.AA.YYYY, --esik 0–100 arasında; --kritik pozitif ve --uyari değerinden küçük; --min-olcum en az 2; --bayat sıfır veya pozitif olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, bugun, esik, a.kritik, a.uyari, a.min_olcum, a.bayat)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(r['kaynaklar'])} kaynak · " + " · ".join(f"{d}: {n}" for d, n in r["say"].items() if n))
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
