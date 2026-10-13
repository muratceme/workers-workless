import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from fractions import Fraction as F
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)


def seri(degerler, kapasite=1000, aralik=1, sistem="S", son=BUGUN):
    n = len(degerler)
    return [{"tarih": son - timedelta(days=aralik * (n - 1 - i)), "sistem": sistem, "kaynak": "D", "kullanilan": D(str(v)), "kapasite": D(kapasite), "birim": "GB"} for i, v in enumerate(degerler)]


def tek(degerler, **kw):
    r = main.raporla(seri(degerler, **{k: v for k, v in kw.items() if k in ("kapasite", "aralik", "son")}), BUGUN, **{k: v for k, v in kw.items() if k not in ("kapasite", "aralik", "son")})
    return r["kaynaklar"][0], [u["tur"] for u in r["uyarilar"]]


class EgilimTesti(unittest.TestCase):
    def test_egim_ve_r2(self):
        self.assertEqual(main.egilim([(0, D(10)), (1, D(12)), (2, D(14))]), (F(2), F(1)))
        self.assertEqual(main.egilim([(0, D(1)), (1, D(3)), (2, D(2))]), (F(1, 2), F(1, 4)))            # sxy = 1, sxx = 2, syy = 2
        self.assertEqual(main.egilim([(0, D(5)), (7, D(5)), (14, D(5))]), (F(0), None))
        self.assertEqual((main.egilim([(0, D(5)), (0, D(9))]), main.egilim([(0, D(9)), (10, D(4))])), (None, (F(-1, 2), F(1))))

    def test_kalan_gun_ve_durum(self):
        k, uy = tek([870, 880, 890])                                                                   # 10 / gün; eşik 900 → 1 gün; dolma 11 gün
        self.assertEqual((k["durum"], k["gunluk"], k["aylik"], k["r2"], k["oran"], k["bos"], k["esik_gun"], k["esik_tarih"], k["dolma_gun"], k["dolma_tarih"], uy),
                         ("Kritik", D("10.000"), D("300.00"), D("1.00"), D("89.00"), D(110), 1, date(2026, 10, 13), 11, date(2026, 10, 23), ["Eşiğe 30 günden az kaldı"]))
        self.assertEqual([(tek(v)[0]["durum"], tek(v)[0]["esik_gun"]) for v in ([804, 807, 810], [801, 804, 807], [624, 627, 630], [621, 624, 627], [0, 3, 7])],
                         [("Kritik", 30), ("Uyarı", 31), ("Uyarı", 90), ("Normal", 91), ("Normal", 256)])  # 3 / gün; son örnek: eğim 3,5 → 893 ÷ 3,5 = 255,1 → 256
        self.assertEqual([(tek(v)[0]["durum"], tek(v)[0]["esik_gun"], tek(v)[0]["dolma_gun"], tek(v)[1]) for v in ([880, 890, 900], [990, 995, 1000], [500, 500, 500], [520, 510, 500], [950, 940, 930], [100, 200])],
                         [("Eşik aşıldı", None, 10, ["Eşik aşıldı"]), ("Eşik aşıldı", None, None, ["Eşik aşıldı"]), ("Büyüme yok", None, None, []), ("Büyüme yok", None, None, []), ("Eşik aşıldı", None, None, ["Eşik aşıldı"]),
                          ("Yetersiz veri", None, None, ["Yetersiz veri"])])
        k, uy = tek([100, 200], min_olcum=2, esik=D(50), kritik=3, uyari_gun=4)                          # 100 / gün; eşik 500 → 3 gün
        self.assertEqual((k["durum"], k["esik_gun"], k["dolma_gun"], uy), ("Kritik", 3, 8, ["Eşiğe 3 günden az kaldı"]))
        k, uy = tek([100, 700, 150, 800, 200, 850])
        self.assertEqual((k["durum"], k["r2"] < D("0.5"), uy), ("Kritik", True, ["Eşiğe 30 günden az kaldı", "Eğilim zayıf"]))

    def test_diger_kurallar(self):
        l = seri([100, 110, 120], son=date(2026, 9, 27)) + seri([10, 20, 30], kapasite=500, sistem="B")
        l[5]["kapasite"] = D(1000)
        l.append({**l[5], "kullanilan": D(40)})                                                          # aynı güne ikinci ölçüm: sonuncusu geçerli
        r = main.raporla(l, BUGUN)
        self.assertEqual([(k["sistem"], k["olcum"], k["son"], k["kapasite"], k["gunluk"], k["esik_gun"]) for k in r["kaynaklar"]], [("B", 3, D(40), D(1000), D("15.000"), 58), ("S", 3, D(120), D(1000), D("10.000"), 78)])
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]), [("Eşiğe 90 günden az kaldı", "B · D"), ("Eşiğe 90 günden az kaldı", "S · D"), ("Kapasite değişmiş", "B · D"), ("Ölçüm eski", "S · D")])
        self.assertEqual(main.raporla(seri([1, 2, 3], son=date(2026, 9, 28)), BUGUN)["uyarilar"], [])
        with self.assertRaises(ValueError):
            main.raporla([], BUGUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "k.xlsx"
            r = main.calistir(ORNEK / "olcumler.csv", cikti, BUGUN)
            self.assertEqual([(k["sistem"], k["durum"], k["oran"], k["gunluk"], k["r2"], k["esik_gun"], k["esik_tarih"], k["dolma_gun"], k["dolma_tarih"]) for k in r["kaynaklar"]],
                             [("MX01", "Eşik aşıldı", D("93.00"), D("1.012"), D("1.00"), None, None, 14, date(2026, 10, 26)),
                              ("FS01", "Kritik", D("83.75"), D("5.000"), D("1.00"), 25, date(2026, 11, 6), 65, date(2026, 12, 16)),
                              ("NAS01", "Uyarı", D("84.38"), D("7.143"), D("1.00"), 63, date(2026, 12, 14), 175, date(2027, 4, 5)),
                              ("DB01", "Normal", D("64.20"), D("0.600"), D("1.00"), 215, date(2027, 5, 15), 299, date(2027, 8, 7)),
                              ("WEB01", "Büyüme yok", D("62.00"), D("-0.004"), D("0.01"), None, None, None, None),
                              ("HV01", "Yetersiz veri", D("61.00"), None, None, None, None, None, None)])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Yüksek", "Eşik aşıldı", "MX01 · C: Posta Kutuları"), ("Yüksek", "Eşiğe 30 günden az kaldı", "FS01 · D: Veri Diski"), ("Orta", "Eşiğe 90 günden az kaldı", "NAS01 · Arşiv"),
                              ("Bilgi", "Yetersiz veri", "HV01 · Sanal Makine Deposu")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Kapasite Özeti", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "o.csv"
            y.write_text("Date;Server;Volume;Used;Total\n2026-10-10;S1;C:;12.5;100\n11.10.2026;S1;C:;1.250,5;2.000,0\n12.10.2026;S1;;15;100\nx;S1;C:;1;100\n12.10.2026;S1;C:;-1;100\n12.10.2026;S1;C:;1;0\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["tarih"], x["kaynak"], x["kullanilan"], x["kapasite"]) for x in l], [u["tur"] for u in uy]),
                             ([(date(2026, 10, 10), "C:", D("12.5"), D(100)), (date(2026, 10, 11), "C:", D("1250.5"), D("2000.0")), (date(2026, 10, 12), "(kaynak yok)", D(15), D(100))], ["Ölçüm okunamadı"] * 3))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--olcumler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "12.10.2026", "--esik", "85", "--kritik", "45", "--uyari", "120", "--min-olcum", "2", "--bayat", "0"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + x) for x in (["--esik", "0"], ["--esik", "x"], ["--kritik", "90"], ["--kritik", "0"], ["--min-olcum", "1"], ["--bayat", "-1"], ["--tarih", "x"])], [2] * 7)
            self.assertEqual(main.main(["--olcumler", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Tarih;Sistem;Kullanılan\n12.10.2026;A;1\n", "Tarih;Sistem;Kullanılan;Kapasite\nx;A;1;1\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
