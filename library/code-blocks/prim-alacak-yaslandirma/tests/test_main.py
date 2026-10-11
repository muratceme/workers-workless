import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
RAPOR = date(2026, 10, 1)
S = [30, 60, 90]


def al(police, gun, kalan, acente="A", brans="Yangın"):
    return {"police": police, "acente": acente, "sigortali": "", "brans": brans, "taksit": "", "vade": RAPOR - timedelta(days=gun), "tutar": D(kalan), "tahsil": D(0), "kalan": D(kalan),
            "gun": max(gun, 0), "aralik": main.aralik_bul(gun, S)}


class YaslandirmaTesti(unittest.TestCase):
    def test_araliklar(self):
        self.assertEqual(main.aralik_adlari(S), ["Vadesi gelmemiş", "1-30", "31-60", "61-90", "90+"])
        self.assertEqual([main.aralik_bul(g, S) for g in (-5, 0, 1, 30, 31, 60, 61, 90, 91, 400)],
                         ["Vadesi gelmemiş", "Vadesi gelmemiş", "1-30", "1-30", "31-60", "31-60", "61-90", "61-90", "90+", "90+"])
        self.assertEqual(main.aralik_adlari([15, 45]), ["Vadesi gelmemiş", "1-15", "16-45", "45+"])

    def test_dagilim(self):
        l = [al("P1", -10, 100), al("P1", 30, 200), al("P2", 45, 300), al("P3", 95, 400, "B", "Kasko"), al("P3", 0, 50, "B", "Kasko")]
        r = main.analiz(l, S)
        g = r["genel"]
        self.assertEqual((g["taksit"], g["police"], g["toplam"], g["gecikmis"], g["oran"], g["gun"]), (5, 3, D(1050), D(900), D("85.71"), 95))
        self.assertEqual(g["aralik"], {"Vadesi gelmemiş": D(150), "1-30": D(200), "31-60": D(300), "61-90": D(0), "90+": D(400)})
        self.assertEqual([(x["ad"], x["police"], x["toplam"], x["gecikmis"], x["oran"]) for x in r["acente"]], [("A", 2, D(600), D(500), D("83.33")), ("B", 1, D(450), D(400), D("88.89"))])
        self.assertEqual([x["ad"] for x in r["brans"]], ["Yangın", "Kasko"])

    def test_police_ve_takip(self):
        l = [al("P1", -10, 100), al("P1", 30, 200), al("P2", 45, 300), al("P2", 15, 50), al("P3", 31, 400), al("P4", -1, 10)]
        r = main.analiz(l, S)
        self.assertEqual([(p["police"], p["gun"], p["gecikmis_taksit"], p["gecikmis"], p["toplam"], p["aralik"]) for p in r["policeler"]],
                         [("P2", 45, 2, D(350), D(350), "31-60"), ("P3", 31, 1, D(400), D(400), "31-60"), ("P1", 30, 1, D(200), D(300), "1-30"), ("P4", 0, 0, D(0), D(10), "Vadesi gelmemiş")])
        self.assertEqual([p["police"] for p in r["takip"]], ["P2", "P3"])                                   # tam 30 gün takip listesine girmez
        self.assertEqual(r["policeler"][0]["en_eski"], RAPOR - timedelta(days=45))
        self.assertEqual([p["police"] for p in main.analiz(l, S, takip_gun=29)["takip"]], ["P2", "P3", "P1"])
        self.assertEqual(main.analiz([], S)["genel"]["toplam"], D(0))


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "y.xlsx"
            r = main.calistir(ORNEK / "prim_taksitleri.csv", cikti, RAPOR, S)
            g = r["genel"]
            self.assertEqual(g["aralik"], {"Vadesi gelmemiş": D(230000), "1-30": D(122000), "31-60": D(60000), "61-90": D(30000), "90+": D(30000)})
            self.assertEqual((g["toplam"], g["gecikmis"], g["oran"], g["gun"], r["sayac"]), (D(472000), D(242000), D("51.27"), 108, {"satir": 14, "kapali": 4}))
            self.assertEqual({x["ad"][-1]: (x["toplam"], x["gecikmis"]) for x in r["acente"]}, {"A": (D(170000), D(120000)), "B": (D(172000), D(92000)), "C": (D(130000), D(30000))})
            self.assertEqual([(p["police"], p["gun"], p["gecikmis"]) for p in r["takip"]], [("P-2", 108, D(90000)), ("P-5", 31, D(30000))])
            self.assertEqual([u["tur"] for u in r["uyarilar"]], ["Fazla tahsilat"])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Yaşlandırma", "Acente Bazında", "Branş Bazında", "Poliçe Bazında", "Takip Listesi", "Açık Taksitler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "t.csv"
            y.write_text("Poliçe No;Vade Tarihi;Taksit Tutarı;Tahsil Edilen\nP1;01.09.2026;1.000,00;250,00\nP2;yarın;100;\nP3;01.09.2026;çok;\nP4;01.09.2026;100;az\nP5;01.09.2026;100;100\n"
                         ";01.09.2026;5;\n", encoding="utf-8")
            l, sayac, uy = main.oku(y, RAPOR, S)
            self.assertEqual([(a["police"], a["kalan"], a["gun"], a["aralik"], a["acente"]) for a in l], [("P1", D("750.00"), 30, "1-30", "Belirtilmemiş"), ("P4", D(100), 30, "1-30", "Belirtilmemiş")])
            self.assertEqual((sayac, sorted(u["tur"] for u in uy)), ({"satir": 3, "kapali": 1}, ["Tahsilat okunamadı", "Taksit tutarı okunamadı", "Vade tarihi okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--taksitler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "01.10.2026", "--araliklar", "15,45", "--takip-gun", "10"]), 0)
            self.assertEqual(main.main(arg + ["--tarih", "dün"]), 2)
            self.assertEqual(main.main(arg + ["--araliklar", "60,30"]), 2)
            self.assertEqual(main.main(["--taksitler", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Poliçe No;Vade Tarihi;Taksit Tutarı\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
