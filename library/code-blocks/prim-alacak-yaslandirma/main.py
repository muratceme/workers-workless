"""
Prim Alacak Yaşlandırma — Workers / Workless kod bloğu
Tahsilat Yönetimi › Tahsilat Uzmanı

Tahsil edilmemiş prim taksitlerini vade gününe göre yaşlandırır; acente ve poliçe bazında raporlar.
  Kalan alacak = taksit tutarı − tahsil edilen            (0 veya altındaysa taksit kapalı sayılır)
  Gecikme günü = rapor tarihi − vade tarihi               (vadesi rapor günü veya sonrası olan taksit "Vadesi gelmemiş")
  Aralıklar    = Vadesi gelmemiş, 1-30, 31-60, 61-90, 90+  (--araliklar ile değişir)
  Takip listesi = en eski açık taksiti --takip-gun (varsayılan 30) günden fazla gecikmiş poliçeler
Kod yalnız listeyi yaşlandırır; ihtar, temerrüt ve poliçe feshi koşullarını değerlendirmez.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"police": ("police no", "police numarasi"), "acente": ("acente", "acente adi"), "sigortali": ("sigortali", "sigortali adi", "musteri"), "brans": ("brans", "urun"),
         "taksit": ("taksit no", "taksit"), "vade": ("vade tarihi", "vade"), "tutar": ("taksit tutari", "tutar", "prim", "tahakkuk"), "tahsil": ("tahsil edilen", "tahsilat", "odenen")}
Z = Decimal(0)


def aralik_adlari(sinirlar: list[int]) -> list[str]:
    return ["Vadesi gelmemiş"] + [f"{a + 1}-{b}" for a, b in zip([0] + sinirlar, sinirlar)] + [f"{sinirlar[-1]}+"]


def aralik_bul(gun: int, sinirlar: list[int]) -> str:
    adlar = aralik_adlari(sinirlar)
    if gun <= 0:
        return adlar[0]
    return next((adlar[i + 1] for i, s in enumerate(sinirlar) if gun <= s), adlar[-1])


def oku(yol: Path, rapor: date, sinirlar: list[int]) -> tuple[list[dict], dict, list[dict]]:
    uy, l = [], []
    sayac = {"satir": 0, "kapali": 0}
    for r in kayitlar(yol, SUTUN, ("police", "vade", "tutar")):
        no = metin(r["police"])
        if not no:
            continue
        kim = f"{no}" + (f" / taksit {metin(r['taksit'])}" if metin(r["taksit"]) else "") + f" (satır {r['_satir']})"
        tutar, tahsil, vade = para(r["tutar"]), para(r["tahsil"]), tarih(r["vade"])
        if tutar is None or tutar < 0:
            uy.append(uyari("Yüksek", "Taksit tutarı okunamadı", kim, f"'{metin(r['tutar'])}'. Satır rapora alınmadı."))
            continue
        if not vade:
            uy.append(uyari("Yüksek", "Vade tarihi okunamadı", kim, f"'{metin(r['vade'])}'. Satır rapora alınmadı."))
            continue
        if tahsil is None and metin(r["tahsil"]):
            uy.append(uyari("Yüksek", "Tahsilat okunamadı", kim, f"'{metin(r['tahsil'])}'. Tahsilat 0 sayıldı."))
        tahsil = tahsil or Z
        sayac["satir"] += 1
        if tahsil > tutar:
            uy.append(uyari("Orta", "Fazla tahsilat", kim, f"Tahsil edilen {tahsil:f} > taksit tutarı {tutar:f}. Fazla kısım ({tahsil - tutar:f}) diğer taksitlere mahsup edilmedi."))
        if tutar - tahsil <= 0:
            sayac["kapali"] += 1
            continue
        gun = (rapor - vade).days
        l.append({"police": no, "acente": metin(r["acente"]) or "Belirtilmemiş", "sigortali": metin(r["sigortali"]), "brans": metin(r["brans"]) or "Belirtilmemiş", "taksit": metin(r["taksit"]),
                  "vade": vade, "tutar": tutar, "tahsil": tahsil, "kalan": tutar - tahsil, "gun": max(gun, 0), "aralik": aralik_bul(gun, sinirlar)})
    return l, sayac, uy


def dagit(alacaklar: list[dict], adlar: list[str], anahtar: str | None = None) -> list[dict]:
    g = {}
    for a in alacaklar:
        x = g.setdefault(a[anahtar] if anahtar else "Toplam", {"ad": a[anahtar] if anahtar else "Toplam", "taksit": 0, "policeler": set(), "toplam": Z, "gecikmis": Z, "gun": 0}
                         | {"aralik": {ad: Z for ad in adlar}})
        x["taksit"] += 1
        x["policeler"].add(a["police"])
        x["toplam"] += a["kalan"]
        x["aralik"][a["aralik"]] += a["kalan"]
        if a["gun"] > 0:
            x["gecikmis"] += a["kalan"]
            x["gun"] = max(x["gun"], a["gun"])
    for x in g.values():
        x["police"] = len(x.pop("policeler"))
        x["oran"] = yuvarla(x["gecikmis"] / x["toplam"] * 100) if x["toplam"] else Z
    return sorted(g.values(), key=lambda x: (-x["gecikmis"], x["ad"]))


def analiz(alacaklar: list[dict], sinirlar: list[int], takip_gun: int = 30) -> dict:
    adlar = aralik_adlari(sinirlar)
    genel = dagit(alacaklar, adlar)[0] if alacaklar else {"ad": "Toplam", "taksit": 0, "police": 0, "toplam": Z, "gecikmis": Z, "gun": 0, "oran": Z, "aralik": {a: Z for a in adlar}}
    pol = {}
    for a in alacaklar:
        p = pol.setdefault(a["police"], {"police": a["police"], "acente": a["acente"], "sigortali": a["sigortali"], "brans": a["brans"], "taksit": 0, "gecikmis_taksit": 0, "toplam": Z,
                                         "gecikmis": Z, "gun": 0, "en_eski": None})
        p["taksit"] += 1
        p["toplam"] += a["kalan"]
        if a["gun"] > 0:
            p["gecikmis_taksit"] += 1
            p["gecikmis"] += a["kalan"]
            if a["gun"] > p["gun"]:
                p["gun"], p["en_eski"] = a["gun"], a["vade"]
    policeler = sorted(pol.values(), key=lambda p: (-p["gun"], -p["gecikmis"], p["police"]))
    for p in policeler:
        p["aralik"] = aralik_bul(p["gun"], sinirlar)
    return {"adlar": adlar, "genel": genel, "acente": dagit(alacaklar, adlar, "acente"), "brans": dagit(alacaklar, adlar, "brans"), "policeler": policeler,
            "takip": [p for p in policeler if p["gun"] > takip_gun], "alacaklar": sorted(alacaklar, key=lambda a: (-a["gun"], a["police"])), "takip_gun": takip_gun}


def rapor_yaz(cikti: Path, r: dict, rapor: date, sayac: dict, uy: list[dict]) -> None:
    wb, adlar, g = Workbook(), r["adlar"], r["genel"]
    renk = lambda ad: None if ad == adlar[0] else KIRMIZI if ad == adlar[-1] else SARI   # noqa: E731
    ws = wb.active
    ws.title = "Yaşlandırma"
    ws.append([f"Rapor tarihi {rapor:%d.%m.%Y} · {sayac['satir']} taksit satırı okundu, {sayac['kapali']} kapalı, {g['taksit']} açık"])
    ws["A1"].font = Font(bold=True)
    _baslik(ws, ["Vade Aralığı (gün)", "Kalan Alacak", "Pay %"], (24, 20, 10), "A3")
    for ad in adlar:
        _satir(ws, [ad, g["aralik"][ad], yuvarla(g["aralik"][ad] / g["toplam"] * 100) if g["toplam"] else Z], para_s=(2, 3), dolgu={1: renk(ad)})
    _satir(ws, ["Toplam açık alacak", g["toplam"], 100], para_s=(2, 3))
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.append([])
    _satir(ws, ["Vadesi geçmiş alacak", g["gecikmis"], g["oran"]], para_s=(2, 3))
    _satir(ws, ["En uzun gecikme (gün)", g["gun"]])
    _satir(ws, [f"Takip listesindeki poliçe (> {r['takip_gun']} gün)", len(r["takip"])])
    for ad, veri, kol in (("Acente Bazında", r["acente"], "Acente"), ("Branş Bazında", r["brans"], "Branş")):
        ws = wb.create_sheet(ad)
        _baslik(ws, [kol, "Poliçe", "Açık Taksit", "Toplam Açık Alacak", *adlar, "Vadesi Geçmiş", "Vadesi Geçmiş %", "En Uzun Gecikme (gün)"], (28, 8, 9, 20, *[16] * len(adlar), 18, 11, 11))
        for x in veri:
            _satir(ws, [x["ad"], x["police"], x["taksit"], x["toplam"], *[x["aralik"][a] or None for a in adlar], x["gecikmis"], x["oran"], x["gun"]],
                   para_s=range(4, 7 + len(adlar)), dolgu={4 + len(adlar): renk(adlar[-1]) if x["aralik"][adlar[-1]] else None})
    ws = wb.create_sheet("Poliçe Bazında")
    _baslik(ws, ["Poliçe No", "Acente", "Sigortalı", "Branş", "Açık Taksit", "Gecikmiş Taksit", "Toplam Açık Alacak", "Vadesi Geçmiş", "En Eski Vade", "Gecikme (gün)", "Aralık"],
            (14, 24, 30, 16, 9, 9, 20, 18, 12, 10, 16))
    for p in r["policeler"]:
        _satir(ws, [p["police"], p["acente"], p["sigortali"], p["brans"], p["taksit"], p["gecikmis_taksit"], p["toplam"], p["gecikmis"], p["en_eski"], p["gun"] or None, p["aralik"]],
               para_s=(7, 8), tarih_s=(9,), dolgu={11: renk(p["aralik"])})
    ws = wb.create_sheet("Takip Listesi")
    _baslik(ws, ["Poliçe No", "Acente", "Sigortalı", "Vadesi Geçmiş", "En Eski Vade", "Gecikme (gün)", "Gecikmiş Taksit", "Yapılan İşlem (siz doldurun)", "Söz Verilen Tarih"],
            (14, 24, 30, 18, 12, 10, 9, 36, 14))
    for p in r["takip"]:
        _satir(ws, [p["police"], p["acente"], p["sigortali"], p["gecikmis"], p["en_eski"], p["gun"], p["gecikmis_taksit"], "", ""], para_s=(4,), tarih_s=(5,))
    ws = wb.create_sheet("Açık Taksitler")
    _baslik(ws, ["Poliçe No", "Taksit No", "Acente", "Sigortalı", "Vade Tarihi", "Taksit Tutarı", "Tahsil Edilen", "Kalan", "Gecikme (gün)", "Aralık"], (14, 8, 24, 30, 12, 16, 16, 16, 10, 16))
    for a in r["alacaklar"]:
        _satir(ws, [a["police"], a["taksit"], a["acente"], a["sigortali"], a["vade"], a["tutar"], a["tahsil"] or None, a["kalan"], a["gun"] or None, a["aralik"]], para_s=(6, 7, 8), tarih_s=(5,),
               dolgu={10: renk(a["aralik"])})
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, rapor: date, sinirlar: list[int], takip_gun: int = 30) -> dict:
    alacaklar, sayac, uy = oku(yol, rapor, sinirlar)
    if not sayac["satir"]:
        raise ValueError(f"{yol.name}: okunabilen taksit satırı yok.")
    r = analiz(alacaklar, sinirlar, takip_gun)
    r["uyarilar"], r["sayac"] = sirala(uy), sayac
    rapor_yaz(cikti, r, rapor, sayac, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Tahsil edilmemiş prim taksitlerini vade gününe göre yaşlandırır.")
    p.add_argument("--taksitler", type=Path, default=ORNEK / "prim_taksitleri.csv", help="Poliçe No, Acente, Sigortalı, Branş, Taksit No, Vade Tarihi, Taksit Tutarı, Tahsil Edilen")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 01.10.2026)")
    p.add_argument("--araliklar", default="30,60,90", help="Aralık üst sınırları, virgülle (varsayılan 30,60,90)")
    p.add_argument("--takip-gun", type=int, default=30, help="Bu günden fazla gecikmiş poliçeler takip listesine alınır (varsayılan 30)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "prim_alacak_yaslandirma.xlsx")
    a = p.parse_args(argv)
    if not a.taksitler.exists():
        print(f"[X] Dosya bulunamadı: {a.taksitler}")
        return 1
    rapor = tarih(a.tarih) if a.tarih else (date(2026, 10, 1) if a.taksitler == ORNEK / "prim_taksitleri.csv" else date.today())
    try:
        sinirlar = [int(x) for x in a.araliklar.split(",")]
        if not rapor or sinirlar != sorted(set(sinirlar)) or sinirlar[0] <= 0 or a.takip_gun < 0:
            raise ValueError
    except ValueError:
        print("[X] --tarih GG.AA.YYYY biçiminde; --araliklar artan pozitif sayılar olmalı (ör. 30,60,90)")
        return 2
    try:
        r = calistir(a.taksitler, a.cikti, rapor, sinirlar, a.takip_gun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    g = r["genel"]
    print(f"[OK] {g['taksit']} açık taksit · {g['police']} poliçe · açık alacak {g['toplam']:,.2f} · vadesi geçmiş {g['gecikmis']:,.2f} (%{g['oran']})")
    for x in r["takip"]:
        print(f"[!] {x['police']} · {x['acente']}: {x['gecikmis']:,.2f} · {x['gun']} gün gecikme")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
