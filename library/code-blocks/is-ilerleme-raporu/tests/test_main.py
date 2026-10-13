import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
R = date(2026, 10, 12)


def k(ad, bas, bit, agirlik, gercek):
    return {"kalem": ad, "bas": date(*bas), "bit": date(*bit), "agirlik": D(agirlik), "gercek": D(gercek)}


PLAN = lambda: [k("Kazı", (2026, 9, 1), (2026, 9, 30), 10, 100), k("Temel", (2026, 9, 21), (2026, 10, 20), 20, 60), k("Kaba", (2026, 10, 11), (2027, 1, 8), 40, 5),   # noqa: E731
                k("Çatı", (2026, 12, 1), (2026, 12, 31), 10, 0), k("Altyapı", (2026, 10, 1), (2026, 10, 10), 20, 80)]


class AnalizTesti(unittest.TestCase):
    def test_kalemler_ve_proje(self):
        r = main.analiz(PLAN(), R)
        self.assertEqual([(x["kalem"], x["sure"], x["pay"], x["plan"], x["sapma"], x["etki"], x["gecikme"], x["durum"]) for x in r["kalemler"]],
                         [("Kazı", 30, D("10.00"), D("100.00"), D("0.00"), D("0.00"), 0, "Tamamlandı"), ("Temel", 30, D("20.00"), D("73.33"), D("-13.33"), D("-2.67"), 4, "Gecikmede"),
                          ("Kaba", 90, D("40.00"), D("2.22"), D("2.78"), D("1.11"), 0, "Planında"), ("Çatı", 31, D("10.00"), D("0.00"), D("0.00"), D("0.00"), 0, "Başlamadı"),
                          ("Altyapı", 10, D("20.00"), D("100.00"), D("-20.00"), D("-4.00"), 2, "Süresi geçti")])
        o = r["ozet"]
        self.assertEqual((o["bas"], o["bit"], o["sure"], o["gecen"], o["plan"], o["gercek"], o["sapma"], o["gecikme"], o["durum"]),
                         (date(2026, 9, 1), date(2027, 1, 8), 130, 42, D("45.56"), D("40.00"), D("-5.56"), 7, "Gecikmede"))
        self.assertEqual([x["kalem"] for x in r["gecikenler"]], ["Altyapı", "Temel"])
        self.assertEqual(r["say"], {"Tamamlandı": 1, "Gecikmede": 1, "Planında": 1, "Başlamadı": 1, "Süresi geçti": 1})

    def test_s_egrisi(self):
        r = main.analiz(PLAN(), R)
        self.assertEqual([(e["tarih"], e["plan"], e["gercek"]) for e in r["egri"]],
                         [(date(2026, 9, 30), D("16.67"), None), (R, D("45.56"), D("40.00")), (date(2026, 10, 31), D("59.33"), None), (date(2026, 11, 30), D("72.67"), None),
                          (date(2026, 12, 31), D("96.44"), None), (date(2027, 1, 8), D("100.00"), None)])
        h = main.analiz([k("A", (2026, 10, 1), (2026, 10, 20), 1, 50)], date(2026, 11, 1), donem="hafta")
        self.assertEqual([(e["tarih"], e["plan"]) for e in h["egri"]], [(date(2026, 10, 7), D("35.00")), (date(2026, 10, 14), D("70.00")), (date(2026, 10, 20), D("100.00"))])
        self.assertEqual((h["ozet"]["gecen"], h["kalemler"][0]["durum"], h["kalemler"][0]["gecikme"]), (20, "Süresi geçti", 10))

    def test_durumlar_ve_tolerans(self):
        d = lambda g, t=R, tol=D(5): main.analiz([k("A", (2026, 10, 3), (2026, 10, 22), 7, g)], t, tolerans=tol)["kalemler"][0]   # noqa: E731  20 gün; 12.10'da 10. gün → %50
        self.assertEqual([(d(g)["plan"], d(g)["durum"]) for g in (0, 44, 45, 55, 56, 100)],
                         [(D("50.00"), "Başlamadı (gecikmede)"), (D("50.00"), "Gecikmede"), (D("50.00"), "Planında"), (D("50.00"), "Planında"), (D("50.00"), "Önde"), (D("50.00"), "Tamamlandı")])
        self.assertEqual((d(10, date(2026, 10, 1))["durum"], d(0, date(2026, 10, 1))["durum"], d(0, date(2026, 10, 1))["plan"], d(99, date(2026, 10, 23))["durum"], d(49, tol=D(0))["durum"]),
                         ("Önde", "Başlamadı", D("0.00"), "Süresi geçti", "Gecikmede"))
        o = main.analiz([k("A", (2026, 10, 3), (2026, 10, 22), 3, 70), k("B", (2026, 10, 3), (2026, 10, 22), 1, 50)], R)["ozet"]
        self.assertEqual((o["plan"], o["gercek"], o["sapma"], o["gecikme"], o["durum"]), (D("50.00"), D("65.00"), D("15.00"), 0, "Önde"))
        with self.assertRaises(ValueError):
            main.analiz([], R)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "i.xlsx"
            r = main.calistir(ORNEK / "is_kalemleri.csv", cikti, R)
            o = r["ozet"]
            self.assertEqual((o["plan"], o["gercek"], o["sapma"], o["durum"], [u["tur"] for u in r["uyarilar"]]), (D("45.56"), D("40.00"), D("-5.56"), "Gecikmede", ["Süresi geçti", "Gecikmede"]))
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Özet", "İş Kalemleri", "S-Eğrisi", "Gecikenler", "Uyarılar"])
            self.assertEqual(len(wb["S-Eğrisi"]._charts), 1)

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "k.csv"
            y.write_text("Faaliyet;Planlanan Başlangıç;Planlanan Bitiş;Bütçe;Tamamlanma %\nA;01.10.2026;20.10.2026;1.250.000,00;%37,5\nB;01.10.2026;20.10.2026;750.000;\n"
                         "C;20.10.2026;01.10.2026;10;5\nD;01.10.2026;20.10.2026;10;120\nE;01.10.2026;20.10.2026;0;5\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["kalem"], x["agirlik"], x["gercek"]) for x in l], [u["tur"] for u in uy]), ([("A", D("1250000.00"), D("37.5")), ("B", D(750000), D(0))], ["Kalem okunamadı"] * 3))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--kalemler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "10.10.2026", "--donem", "hafta"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(arg + ["--tarih", "dün"]), 2)
            self.assertEqual(main.main(arg + ["--tolerans", "-1"]), 2)
            self.assertEqual(main.main(["--kalemler", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("İş Kalemi;Başlangıç\nA;1\n", "İş Kalemi;Başlangıç;Bitiş;Ağırlık;Gerçekleşen\nA;x;y;1;1\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
