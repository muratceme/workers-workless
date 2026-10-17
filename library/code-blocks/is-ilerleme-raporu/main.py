"""
İş İlerleme Raporu — Gerisi Bende kod bloğu
Şantiye › Şantiye Şefi

İş programı ve gerçekleşen imalat yüzdelerinden planlanan-gerçekleşen ilerlemeyi, S-eğrisini ve geciken iş kalemlerini raporlar.
  Kalem planlanan %  = (rapor tarihi − başlangıç + 1) ÷ (bitiş − başlangıç + 1) × 100, 0–100 arasında        (takvim günü, iş süreye eşit dağılmış varsayılır)
  Kalem payı %       = kalemin ağırlığı (veya bütçesi) ÷ ağırlıklar toplamı × 100
  Proje ilerlemesi   = Σ pay × kalem yüzdesi ÷ 100                                                         (planlanan ve gerçekleşen ayrı ayrı)
  Sapma (puan)       = gerçekleşen % − planlanan %;   projeye etkisi = pay × sapma ÷ 100
  Gecikme (gün, kaba)= (planlanan % − gerçekleşen %) ÷ 100 × kalemin süresi
  Durum              = Tamamlandı · Süresi geçti (bitiş geçmiş, iş bitmemiş) · Başlamadı (gecikmede) · Gecikmede (sapma < −tolerans) · Önde · Planında · Başlamadı
  S-eğrisi           = ay (veya hafta) sonlarında planlanan kümülatif proje ilerlemesi; rapor tarihinde gerçekleşen değerle birlikte
Doğrusal dağılım bir varsayımdır: işin başı veya sonu yoğun kalemlerde planlanan yüzde gerçek programdan farklı çıkar. Kritik yol hesabı yapılmaz.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
from openpyxl.chart import LineChart, Reference

SUTUN = {"kalem": ("is kalemi", "kalem", "faaliyet", "imalat", "is"), "bas": ("baslangic", "baslangic tarihi", "planlanan baslangic"), "bit": ("bitis", "bitis tarihi", "planlanan bitis"),
         "agirlik": ("agirlik", "butce", "tutar", "sozlesme tutari", "pursantaj"), "gercek": ("gerceklesen", "gerceklesen ilerleme", "tamamlanma", "fiili", "ilerleme")}
Z, YUZ = Decimal(0), Decimal(100)
DURUM_RENK = {"Süresi geçti": KIRMIZI, "Başlamadı (gecikmede)": KIRMIZI, "Gecikmede": KIRMIZI, "Önde": YESIL, "Tamamlandı": YESIL}


def olcu(x) -> Decimal | None:
    """Yüzde ve ölçü: 12,5 = 12.5; 1.250,5 = bin iki yüz elli buçuk."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    else:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("kalem", "bas", "bit", "agirlik", "gercek")):
        ad = metin(r["kalem"])
        if not ad:
            continue
        b, s, a, g = tarih(r["bas"]), tarih(r["bit"]), para(r["agirlik"]), olcu(r["gercek"]) if metin(r["gercek"]) else Z
        if b is None or s is None or s < b or a is None or a <= 0 or g is None or g < 0 or g > 100:
            uy.append(uyari("Yüksek", "Kalem okunamadı", f"{ad} (satır {r['_satir']})", f"Başlangıç '{metin(r['bas'])}', bitiş '{metin(r['bit'])}', ağırlık '{metin(r['agirlik'])}', gerçekleşen '{metin(r['gercek'])}'. "
                            "Bitiş başlangıçtan önce olamaz, ağırlık pozitif, gerçekleşen 0–100 arasında olmalı. Satır alınmadı; proje yüzdeleri bu kalem olmadan hesaplandı."))
            continue
        l.append({"kalem": ad, "bas": b, "bit": s, "agirlik": a, "gercek": g})
    return l, uy


def planlanan(k: dict, t: date) -> Decimal:
    sure = (k["bit"] - k["bas"]).days + 1
    return Decimal(min(max((t - k["bas"]).days + 1, 0), sure)) / sure * 100


def analiz(kalemler: list[dict], rapor: date, tolerans: Decimal = Decimal(5), donem: str = "ay") -> dict:
    if not kalemler:
        raise ValueError("Değerlendirilecek iş kalemi yok.")
    top = sum(k["agirlik"] for k in kalemler)
    bas, bit = min(k["bas"] for k in kalemler), max(k["bit"] for k in kalemler)
    liste, plan_p, gercek_p = [], Z, Z
    for k in kalemler:
        sure, pl, g, pay = (k["bit"] - k["bas"]).days + 1, planlanan(k, rapor), k["gercek"], k["agirlik"] / top
        sapma = g - pl
        if g >= 100:
            durum = "Tamamlandı"
        elif rapor < k["bas"]:
            durum = "Önde" if g > 0 else "Başlamadı"
        elif rapor > k["bit"]:
            durum = "Süresi geçti"
        elif g == 0:
            durum = "Başlamadı (gecikmede)"
        else:
            durum = "Gecikmede" if sapma < -tolerans else "Önde" if sapma > tolerans else "Planında"
        plan_p += pay * pl
        gercek_p += pay * g
        liste.append({"kalem": k["kalem"], "bas": k["bas"], "bit": k["bit"], "sure": sure, "pay": yuvarla(pay * 100), "plan": yuvarla(pl), "gercek": yuvarla(g), "sapma": yuvarla(sapma), "etki": yuvarla(pay * sapma),
                      "gecikme": int(yuvarla(-sapma / 100 * sure, "1")) if sapma < 0 else 0, "durum": durum})
    sure_p = (bit - bas).days + 1
    sapma_p = gercek_p - plan_p
    ozet = {"rapor": rapor, "bas": bas, "bit": bit, "sure": sure_p, "gecen": min(max((rapor - bas).days + 1, 0), sure_p), "plan": yuvarla(plan_p), "gercek": yuvarla(gercek_p), "sapma": yuvarla(sapma_p),
            "gecikme": int(yuvarla(-sapma_p / 100 * sure_p, "1")) if sapma_p < 0 else 0,
            "durum": "Gecikmede" if sapma_p < -tolerans else "Önde" if sapma_p > tolerans else "Planında", "tolerans": tolerans}
    noktalar, t = set(), bas
    while t <= bit:
        if donem == "hafta":
            t += timedelta(days=7)
            noktalar.add(min(t - timedelta(days=1), bit))
        else:
            t = (t.replace(day=1) + timedelta(days=32)).replace(day=1)
            noktalar.add(min(t - timedelta(days=1), bit))
    if bas <= rapor <= bit:
        noktalar.add(rapor)
    egri = [{"tarih": n, "plan": yuvarla(sum(k["agirlik"] / top * planlanan(k, n) for k in kalemler)), "gercek": yuvarla(gercek_p) if n == rapor else None} for n in sorted(noktalar)]
    gecikenler = sorted((x for x in liste if x["etki"] < 0), key=lambda x: (x["etki"], x["kalem"]))
    say = defaultdict(int)
    for x in liste:
        say[x["durum"]] += 1
    return {"kalemler": liste, "ozet": ozet, "egri": egri, "gecikenler": gecikenler, "say": dict(say), "uyarilar": []}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb, o = Workbook(), r["ozet"]
    ws = wb.active
    ws.title = "Özet"
    _baslik(ws, ["Gösterge", "Değer", "Açıklama"], (30, 16, 90))
    for ad, deger, ac, bicim in (("Rapor tarihi", o["rapor"], "", TF), ("Proje başlangıcı", o["bas"], "En erken kalem başlangıcı", TF), ("Proje bitişi", o["bit"], "En geç kalem bitişi", TF),
                                 ("Proje süresi (gün)", o["sure"], "Takvim günü", None), ("Geçen süre (gün)", o["gecen"], "", None), ("Planlanan ilerleme %", o["plan"], "Σ pay × kalemin planlanan yüzdesi", PF),
                                 ("Gerçekleşen ilerleme %", o["gercek"], "Σ pay × kalemin gerçekleşen yüzdesi", PF), ("Sapma (puan)", o["sapma"], "Gerçekleşen − planlanan", PF),
                                 ("Tahmini gecikme (gün, kaba)", o["gecikme"], "Sapma ÷ 100 × proje süresi; kritik yol dikkate alınmaz", None),
                                 ("Durum", o["durum"], f"Sapma ±{o['tolerans'].normalize():f} puanı aşarsa 'Gecikmede' veya 'Önde'", None)):
        n = _satir(ws, [ad, deger, ac], sar=(3,))
        if bicim:
            ws.cell(n, 2).number_format = bicim
    ws.cell(ws.max_row, 2).fill = PatternFill("solid", fgColor=KIRMIZI if o["durum"] == "Gecikmede" else YESIL)
    ws = wb.create_sheet("İş Kalemleri")
    _baslik(ws, ["İş Kalemi", "Başlangıç", "Bitiş", "Süre (gün)", "Pay %", "Planlanan %", "Gerçekleşen %", "Sapma (puan)", "Projeye Etkisi (puan)", "Gecikme (gün)", "Durum", "Not (siz doldurun)"],
            (34, 12, 12, 9, 9, 11, 12, 10, 12, 10, 22, 30))
    for x in r["kalemler"]:
        _satir(ws, [x["kalem"], x["bas"], x["bit"], x["sure"], x["pay"], x["plan"], x["gercek"], x["sapma"], x["etki"], x["gecikme"] or None, x["durum"], ""], para_s=(5, 6, 7, 8, 9), tarih_s=(2, 3),
               dolgu={11: DURUM_RENK.get(x["durum"])})
    ws = wb.create_sheet("S-Eğrisi")
    _baslik(ws, ["Tarih", "Planlanan Kümülatif %", "Gerçekleşen %"], (14, 20, 16))
    for e in r["egri"]:
        _satir(ws, [e["tarih"], e["plan"], e["gercek"]], para_s=(2, 3), tarih_s=(1,), dolgu={3: SARI if e["gercek"] is not None else None})
    if len(r["egri"]) > 1:
        g = LineChart()
        g.title, g.height, g.width = "Planlanan ilerleme (S-eğrisi)", 9, 20
        g.y_axis.title, g.y_axis.scaling.min, g.y_axis.scaling.max = "%", 0, 100
        g.add_data(Reference(ws, min_col=2, min_row=1, max_row=ws.max_row), titles_from_data=True)
        g.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
        ws.add_chart(g, "E2")
    ws = wb.create_sheet("Gecikenler")
    _baslik(ws, ["İş Kalemi", "Durum", "Planlanan %", "Gerçekleşen %", "Sapma (puan)", "Projeye Etkisi (puan)", "Gecikme (gün)", "Bitiş"], (34, 22, 11, 12, 10, 12, 10, 12))
    for x in r["gecikenler"]:
        _satir(ws, [x["kalem"], x["durum"], x["plan"], x["gercek"], x["sapma"], x["etki"], x["gecikme"] or None, x["bit"]], para_s=(3, 4, 5, 6), tarih_s=(8,))
    ws.append([])
    ws.append(["Proje ilerlemesini en çok geri çeken kalemler üstte. Bu sıralama kritik yol değildir: bolluğu olan bir kalemin gecikmesi proje bitişini etkilemeyebilir."])
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, rapor: date, **ayar) -> dict:
    kalemler, uy = oku(yol)
    r = analiz(kalemler, rapor, **ayar)
    for x in r["kalemler"]:
        if x["durum"] in ("Süresi geçti", "Başlamadı (gecikmede)"):
            uy.append(uyari("Yüksek", x["durum"], x["kalem"], f"Planlanan %{x['plan']}, gerçekleşen %{x['gercek']}; plan {x['bas']:%d.%m.%Y} – {x['bit']:%d.%m.%Y}."))
        elif x["durum"] == "Gecikmede":
            uy.append(uyari("Orta", "Gecikmede", x["kalem"], f"Planlanan %{x['plan']}, gerçekleşen %{x['gercek']} (sapma {x['sapma']} puan, yaklaşık {x['gecikme']} gün)."))
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="İş programı ve gerçekleşen yüzdelerden ilerleme raporu ve S-eğrisi üretir.")
    p.add_argument("--kalemler", type=Path, default=ORNEK / "is_kalemleri.csv", help="İş Kalemi, Başlangıç, Bitiş, Ağırlık (veya Bütçe), Gerçekleşen %%")
    p.add_argument("--tarih", help="Rapor tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--tolerans", default="5", help="Sapma bu puanı aşarsa 'Gecikmede' / 'Önde' (varsayılan 5)")
    p.add_argument("--donem", choices=("ay", "hafta"), default="ay", help="S-eğrisi dönemi (varsayılan ay)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "is_ilerleme_raporu.xlsx")
    a = p.parse_args(argv)
    ornek = a.kalemler == ORNEK / "is_kalemleri.csv"
    rapor = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    tol = olcu(a.tolerans)
    if rapor is None or tol is None or tol < 0:
        print("[X] --tarih GG.AA.YYYY biçiminde, --tolerans sıfır veya pozitif sayı olmalı")
        return 2
    if not a.kalemler.exists():
        print(f"[X] Dosya bulunamadı: {a.kalemler}")
        return 1
    try:
        r = calistir(a.kalemler, a.cikti, rapor, tolerans=tol, donem=a.donem)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {rapor:%d.%m.%Y} · {len(r['kalemler'])} iş kalemi · planlanan %{o['plan']} · gerçekleşen %{o['gercek']} · sapma {o['sapma']} puan · {o['durum']}")
    for x in r["gecikenler"][:5]:
        print(f"[!] {x['kalem']}: {x['durum']}, sapma {x['sapma']} puan (projeye etkisi {x['etki']})")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
