import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
DONEM = (date(2026, 1, 1), date(2026, 12, 31))


def b(kanal, kampanya, butce):
    return {"kanal": kanal, "kampanya": kampanya, "butce": D(str(butce))}


def h(kanal, kampanya, harcama, gelir=None, donusum=None, t=date(2026, 3, 1)):
    d = lambda x: None if x is None else D(str(x))   # noqa: E731
    return {"tarih": t, "kanal": kanal, "kampanya": kampanya, "harcama": D(str(harcama)), "gelir": d(gelir), "donusum": d(donusum)}


class HesapTesti(unittest.TestCase):
    def test_olcum(self):
        r = main.hesapla([b("A", "x", 1000), b("A", "y", 500), b("B", "z", 0)], [h("A", "x", 400, 1200, 8), h("a", "X", 200), h("A", "x", 100, 0, 0), h("A", "y", 600, 300), h("C", "w", 50, t=None), h("B", "z", 0)])
        self.assertEqual([(o["kanal"], o["kampanya"], o["butce"], o["harcama"], o["kalan"], o["kullanim"], o["gelir"], o["roas"], o["roi"], o["donusum"], o["dbm"], o["hiz"]) for o in r["kampanyalar"]],
                         [("A", "x", D(1000), D(700), D(300), D("70.00"), D(1200), D("2.40"), D("140.00"), D(8), D("62.50"), None), ("A", "y", D(500), D(600), D(-100), D("120.00"), D(300), D("0.50"), D("-50.00"), None, None, "Aşım"),
                          ("B", "z", D(0), D(0), D(0), None, None, None, None, None, None, None), ("C", "w", None, D(50), None, None, None, None, None, None, None, None)])
        self.assertEqual([(o["kanal"], o["butce"], o["harcama"], o["kalan"], o["kullanim"], o["gelir"], o["roas"]) for o in r["kanallar"]],
                         [("A", D(1500), D(1300), D(200), D("86.67"), D(1500), D("1.36")), ("B", D(0), D(0), D(0), None, None, None), ("C", D(0), D(50), D(-50), None, None, None)])
        t = r["toplam"]
        self.assertEqual((t["butce"], t["harcama"], t["kalan"], t["kullanim"], t["gelir"], t["roas"], t["roi"], t["dbm"], r["beklenen"]), (D(1500), D(1350), D(150), D("90.00"), D(1500), D("1.36"), D("36.36"), D("62.50"), None))
        self.assertEqual([(o["ay"], o["harcama"]) for o in r["aylar"]], [("(tarih yok)", D(50)), ("2026-03", D(1300))])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Bilgi", "Gelir yazılmamış", "C › w"), ("Orta", "Bütçesiz harcama", "C › w"), ("Orta", "ROAS düşük", "A › y"), ("Yüksek", "Bütçe aşımı", "A › y")])
        self.assertEqual([u["tur"] for u in main.hesapla([b("A", "x", 100)], [h("A", "x", 10, 19)], roas_esik=D(2))["uyarilar"]], ["ROAS düşük"])
        self.assertEqual(main.hesapla([b("A", "x", 100)], [h("A", "x", 10, 20)], roas_esik=D(2))["uyarilar"], [])
        r = main.hesapla([b("A", "x", 100), b("a", "X", 50)], [])
        self.assertEqual((r["kampanyalar"][0]["butce"], [u["tur"] for u in r["uyarilar"]], r["aylar"]), (D(150), ["Bütçe satırı tekrarlanıyor"], []))
        with self.assertRaises(ValueError):
            main.hesapla([], [])

    def test_hiz(self):
        harcamalar = (424, "423.9", 124, "124.1", "1000.01", 1000)
        r = main.hesapla([b("K", str(i), 1000) for i in range(6)], [h("K", str(i), x) for i, x in enumerate(harcamalar)], DONEM, date(2026, 4, 10))
        self.assertEqual(r["beklenen"], D("27.40"))
        self.assertEqual([(o["kullanim"], o["fark"], o["hiz"], o["tahmin"]) for o in r["kampanyalar"]],
                         [(D("42.40"), D("15.00"), "Hızlı", D("1547.60")), (D("42.39"), D("14.99"), "Planında", D("1547.24")), (D("12.40"), D("-15.00"), "Yavaş", D("452.60")), (D("12.41"), D("-14.99"), "Planında", D("452.97")),
                          (D("100.00"), D("72.60"), "Aşım", D("3650.04")), (D("100.00"), D("72.60"), "Hızlı", D("3650.00"))])
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"] if u["tur"] != "Gelir yazılmamış"),
                         [("Bütçe aşımı", "K › 4"), ("Harcama hızı düşük", "K › 2"), ("Harcama hızı yüksek", "K › 0"), ("Harcama hızı yüksek", "K › 5")])
        hz = lambda bugun, x, **kw: [(o["hiz"], o["tahmin"]) for o in main.hesapla([b("K", "k", 1000)], [h("K", "k", x)], DONEM, bugun, **kw)["kampanyalar"]] + [   # noqa: E731
            main.hesapla([b("K", "k", 1000)], [h("K", "k", x)], DONEM, bugun, **kw)["beklenen"]]
        self.assertEqual((hz(date(2025, 12, 1), 0), hz(date(2025, 12, 1), 150), hz(date(2027, 3, 1), 850), hz(date(2026, 4, 10), 380, hiz=D(10)), hz(date(2026, 1, 1), 0)),
                         ([("Planında", None), D("0.00")], [("Hızlı", None), D("0.00")], [("Yavaş", D("850.00")), D("100.00")], [("Hızlı", D("1387.00")), D("27.40")], [("Planında", D("0.00")), D("0.27")]))
        r = main.hesapla([b("K", "k", 1000)], [h("K", "k", 100, t=date(2025, 12, 31)), h("K", "k", 100, t=date(2026, 1, 1)), h("K", "k", 30, t=None)], DONEM, date(2026, 4, 10))
        self.assertEqual([(u["onem"], u["kim"]) for u in r["uyarilar"] if u["tur"] == "Dönem dışı harcama"], [("Orta", "1 satır")])
        self.assertEqual(r["kampanyalar"][0]["harcama"], D(230))


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "butce.csv", ORNEK / "harcamalar.csv", cikti, DONEM, date(2026, 10, 12))
            self.assertEqual(r["beklenen"], D("78.08"))
            self.assertEqual([(o["kampanya"], o["harcama"], o["kullanim"], o["roas"], o["hiz"]) for o in r["kampanyalar"]],
                             [("Marka Aramaları", D(190000), D("79.17"), D("6.00"), "Planında"), ("Ürün Aramaları", D(585000), D("97.50"), D("3.00"), "Hızlı"), ("Aylık Bülten", D(45000), D("75.00"), D("7.00"), "Planında"),
                              ("Bayi Toplantısı", D(0), D("0.00"), None, "Yavaş"), ("Sektör Fuarı", D(450000), D("100.00"), None, "Hızlı"), ("Bahar Kampanyası", D(312000), D("104.00"), D("0.90"), "Aşım"),
                              ("Influencer İş Birliği", D(40000), None, D("1.30"), None), ("Yeniden Hedefleme", D(90000), D("50.00"), D("4.50"), "Yavaş")])
            o = r["kampanyalar"][1]
            self.assertEqual((o["kalan"], o["roi"], o["donusum"], o["dbm"], o["fark"], o["tahmin"]), (D(15000), D("200.00"), D(1300), D("450.00"), D("19.42"), D("749210.53")))
            self.assertEqual([(o["kanal"], o["butce"], o["harcama"], o["kalan"], o["kullanim"], o["gelir"], o["roas"], o["roi"], o["donusum"], o["dbm"]) for o in r["kanallar"]],
                             [("Arama Reklamı", D(840000), D(775000), D(65000), D("92.26"), D(2895000), D("3.74"), D("273.55"), D(2250), D("344.44")), ("E-posta", D(60000), D(45000), D(15000), D("75.00"), D(315000), D("7.00"), D("600.00"), D(420), D("107.14")),
                              ("Etkinlik", D(200000), D(0), D(200000), D("0.00"), None, None, None, None, None), ("Fuar", D(450000), D(450000), D(0), D("100.00"), None, None, None, D(85), D("5294.12")),
                              ("Sosyal Medya", D(480000), D(442000), D(38000), D("92.08"), D(737800), D("1.67"), D("66.92"), D(880), D("502.27"))])
            t = r["toplam"]
            self.assertEqual((t["butce"], t["harcama"], t["kalan"], t["kullanim"], t["gelir"], t["roas"], t["roi"], t["donusum"], t["dbm"]),
                             (D(2030000), D(1712000), D(318000), D("84.33"), D(3947800), D("3.13"), D("212.82"), D(3635), D("470.98")))
            self.assertEqual([(o["ay"], o["harcama"], o["gelir"]) for o in r["aylar"]][:3], [("2026-02", D(15000), D(105000)), ("2026-03", D(270000), D(1020000)), ("2026-04", D(180000), D(162000))])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Yüksek", "Bütçe aşımı", "Sosyal Medya › Bahar Kampanyası"), ("Orta", "Bütçesiz harcama", "Sosyal Medya › Influencer İş Birliği"), ("Orta", "Harcama hızı yüksek", "Arama Reklamı › Ürün Aramaları"),
                              ("Orta", "Harcama hızı yüksek", "Fuar › Sektör Fuarı"), ("Orta", "ROAS düşük", "Sosyal Medya › Bahar Kampanyası"), ("Bilgi", "Gelir yazılmamış", "Fuar › Sektör Fuarı"),
                              ("Bilgi", "Harcama hızı düşük", "Etkinlik › Bayi Toplantısı"), ("Bilgi", "Harcama hızı düşük", "Sosyal Medya › Yeniden Hedefleme")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Kampanyalar", "Kanallar", "Aylar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            by, hy = Path(tmp) / "b.csv", Path(tmp) / "h.csv"
            by.write_text("Mecra;Faaliyet;Planlanan\nTV;;1.250.000\nRadyo;Sabah;x\n;;5\n", encoding="utf-8")
            hy.write_text("Harcama Tarihi;Mecra;Faaliyet;Tutar;Ciro;Satış Adedi\n05.03.2026;TV;;250.000,50;;12\nx;TV;;-1;;\n05.03.2026;TV;;5;-1;\n;;;7;;\n", encoding="utf-8")
            l, uy = main.oku_butce(by)
            self.assertEqual((l, [(u["tur"], u["kim"]) for u in uy]), ([{"kanal": "TV", "kampanya": "(genel)", "butce": D(1250000)}], [("Bütçe okunamadı", "Radyo › Sabah")]))
            l, uy = main.oku_harcama(hy)
            self.assertEqual((l, [(u["tur"], u["kim"]) for u in uy]),
                             ([{"tarih": date(2026, 3, 5), "kanal": "TV", "kampanya": "(genel)", "harcama": D("250000.50"), "gelir": None, "donusum": D(12)}], [("Harcama satırı okunamadı", "TV")] * 2))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--butce", str(by), "--harcamalar", str(hy), "--cikti", c]
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(arg + ["--donem-baslangic", "01.01.2026", "--donem-bitis", "31.12.2026", "--tarih", "01.07.2026", "--hiz", "10", "--roas", "2,5"]), 0)
            self.assertEqual([main.main(arg + x) for x in (["--donem-baslangic", "01.01.2026"], ["--donem-baslangic", "31.12.2026", "--donem-bitis", "01.01.2026"], ["--donem-baslangic", "x", "--donem-bitis", "01.01.2026"],
                                                            ["--tarih", "x"], ["--hiz", "0"], ["--roas=-1"])] + [main.main(["--butce", str(by)])], [2] * 7)
            self.assertEqual(main.main(["--butce", str(by), "--harcamalar", str(Path(tmp) / "yok.csv")]), 1)
            hy.write_text("Kanal;Kampanya\nA;1\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            by.write_text("Kanal;Bütçe\n;\n", encoding="utf-8")
            hy.write_text("Kanal;Harcama\n;\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
