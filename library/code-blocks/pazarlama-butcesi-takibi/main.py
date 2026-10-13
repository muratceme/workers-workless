"""
Pazarlama Bütçesi Takibi — Workers / Workless kod bloğu
Pazarlama › Pazarlama Müdürü

Kanal ve kampanya bazında pazarlama bütçesini, harcamayı ve getiriyi izler.
  Kalan bütçe          = bütçe − harcama;   kullanım % = harcama ÷ bütçe × 100
  ROAS                 = gelir ÷ geliri yazılmış harcama   (reklam harcamasının getirisi; 1'in altı harcamayı karşılamıyor)
  ROI %                = (gelir − geliri yazılmış harcama) ÷ geliri yazılmış harcama × 100
  Dönüşüm başı maliyet = dönüşümü yazılmış harcama ÷ dönüşüm sayısı
  Beklenen kullanım %  = dönemin geçen gün sayısı ÷ dönemin toplam gün sayısı × 100   (dönem verilirse)
  Hız                  = kullanım % − beklenen kullanım %; fark eşiği (15 puan) aşarsa "Hızlı" ya da "Yavaş"
  Dönem sonu tahmini   = harcama × toplam gün ÷ geçen gün   (harcama aynı hızla sürerse)
Gelir, sizin kampanyaya atfettiğiniz tutardır; kod atıf yapmaz. ROI brüt gelir üzerindendir, ürün maliyeti düşülmez.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
B_SUTUN = {"kanal": ("kanal", "mecra"), "kampanya": ("kampanya", "kampanya adi", "faaliyet"), "butce": ("butce", "butce tutari", "planlanan")}
H_SUTUN = {"tarih": ("tarih", "harcama tarihi"), "kanal": ("kanal", "mecra"), "kampanya": ("kampanya", "kampanya adi", "faaliyet"), "harcama": ("harcama", "tutar", "maliyet"), "gelir": ("gelir", "ciro", "satis"),
           "donusum": ("donusum", "donusum sayisi", "satis adedi", "potansiyel musteri")}
Z, YUZ = Decimal(0), Decimal(100)
HIZ_RENK = {"Aşım": KIRMIZI, "Hızlı": SARI, "Yavaş": MAVI, "Planında": YESIL}


def oku_butce(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, B_SUTUN, ("kanal", "butce")):
        if not metin(r["kanal"]):
            continue
        b = para(r["butce"])
        if b is None or b < 0:
            uy.append(uyari("Yüksek", "Bütçe okunamadı", f"{metin(r['kanal'])} › {metin(r['kampanya'])}".strip(" ›"), f"Bütçe '{metin(r['butce'])}' okunamadı (satır {r['_satir']}). Satır alınmadı."))
            continue
        l.append({"kanal": metin(r["kanal"]), "kampanya": metin(r["kampanya"]) or "(genel)", "butce": b})
    return l, uy


def oku_harcama(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, H_SUTUN, ("kanal", "harcama")):
        if not metin(r["kanal"]):
            continue
        h, g, d = para(r["harcama"]), para(r["gelir"]), para(r["donusum"])
        if h is None or h < 0 or (g is not None and g < 0) or (d is not None and d < 0):
            uy.append(uyari("Yüksek", "Harcama satırı okunamadı", f"{metin(r['kanal'])} › {metin(r['kampanya'])}".strip(" ›"), f"Satır {r['_satir']}: harcama '{metin(r['harcama'])}', gelir '{metin(r['gelir'])}', dönüşüm '{metin(r['donusum'])}'. "
                            "Eksi veya okunamayan değer var; satır alınmadı."))
            continue
        l.append({"tarih": tarih(r["tarih"]), "kanal": metin(r["kanal"]), "kampanya": metin(r["kampanya"]) or "(genel)", "harcama": h, "gelir": g, "donusum": d})
    return l, uy


def _olc(satirlar: list[dict], butce: Decimal | None) -> dict:
    h = sum((s["harcama"] for s in satirlar), Z)
    gelirli, donusumlu = [s for s in satirlar if s["gelir"] is not None], [s for s in satirlar if s["donusum"] is not None]
    g, gh = sum((s["gelir"] for s in gelirli), Z), sum((s["harcama"] for s in gelirli), Z)
    d, dh = sum((s["donusum"] for s in donusumlu), Z), sum((s["harcama"] for s in donusumlu), Z)
    return {"butce": butce, "harcama": h, "kalan": butce - h if butce is not None else None, "kullanim": yuvarla(h * YUZ / butce) if butce else None, "gelir": g if gelirli else None,
            "roas": yuvarla(g / gh) if gh else None, "roi": yuvarla((g - gh) * YUZ / gh) if gh else None, "donusum": d if donusumlu else None, "dbm": yuvarla(dh / d) if d else None}


def hesapla(butce: list[dict], hareket: list[dict], donem: tuple[date, date] | None = None, bugun: date | None = None, hiz: Decimal = Decimal(15), roas_esik: Decimal = Decimal(1)) -> dict:
    if not butce and not hareket:
        raise ValueError("Bütçe ve harcama listesi boş.")
    uy, kampanyalar, sira = [], {}, []
    for b in butce:
        k = (katla(b["kanal"]), katla(b["kampanya"]))
        if k in kampanyalar:
            kampanyalar[k]["butce"] += b["butce"]
            uy.append(uyari("Bilgi", "Bütçe satırı tekrarlanıyor", f"{b['kanal']} › {b['kampanya']}", "Aynı kanal ve kampanya bütçe dosyasında birden çok satırda; tutarlar toplandı."))
        else:
            kampanyalar[k] = {"kanal": b["kanal"], "kampanya": b["kampanya"], "butce": b["butce"], "satirlar": []}
    for h in hareket:
        k = (katla(h["kanal"]), katla(h["kampanya"]))
        if k not in kampanyalar:
            kampanyalar[k] = {"kanal": h["kanal"], "kampanya": h["kampanya"], "butce": None, "satirlar": []}
        kampanyalar[k]["satirlar"].append(h)
    beklenen = toplam_gun = gecen = None
    if donem and bugun:
        toplam_gun = (donem[1] - donem[0]).days + 1
        gecen = min(max((bugun - donem[0]).days + 1, 0), toplam_gun)
        beklenen = yuvarla(Decimal(gecen * 100) / toplam_gun)
        disarida = [h for h in hareket if h["tarih"] and not donem[0] <= h["tarih"] <= donem[1]]
        if disarida:
            uy.append(uyari("Orta", "Dönem dışı harcama", f"{len(disarida)} satır", f"Tarihi {donem[0]:%d.%m.%Y} – {donem[1]:%d.%m.%Y} dışında kalan harcama toplamı {sum((h['harcama'] for h in disarida), Z):,.2f}. Hesaba dahil edildi; dosyayı kontrol edin."))
    liste = []
    for c in kampanyalar.values():
        o = {"kanal": c["kanal"], "kampanya": c["kampanya"], **_olc(c["satirlar"], c["butce"]), "hiz": None, "fark": None, "tahmin": None}
        kim = f"{o['kanal']} › {o['kampanya']}"
        if o["butce"] is None:
            uy.append(uyari("Orta", "Bütçesiz harcama", kim, f"Bütçe dosyasında olmayan kampanyaya {o['harcama']:,.2f} harcanmış."))
        elif o["harcama"] > o["butce"]:
            o["hiz"] = "Aşım"
            uy.append(uyari("Yüksek", "Bütçe aşımı", kim, f"Bütçe {o['butce']:,.2f}, harcama {o['harcama']:,.2f}; aşım {-o['kalan']:,.2f}" + (f" (kullanım %{o['kullanim']})." if o["kullanim"] is not None else ".")))
        if beklenen is not None and o["kullanim"] is not None:
            o["fark"] = o["kullanim"] - beklenen
            o["tahmin"] = yuvarla(o["harcama"] * toplam_gun / gecen) if gecen else None
            if o["hiz"] is None:
                o["hiz"] = "Hızlı" if o["fark"] >= hiz else "Yavaş" if o["fark"] <= -hiz else "Planında"
                if o["hiz"] == "Hızlı":
                    uy.append(uyari("Orta", "Harcama hızı yüksek", kim, f"Dönemin %{beklenen}'i geçti, bütçenin %{o['kullanim']}'i harcandı." + (f" Aynı hızla dönem sonu tahmini {o['tahmin']:,.2f} (bütçe {o['butce']:,.2f})." if o["tahmin"] is not None else "")))
                elif o["hiz"] == "Yavaş":
                    uy.append(uyari("Bilgi", "Harcama hızı düşük", kim, f"Dönemin %{beklenen}'i geçti, bütçenin %{o['kullanim']}'i harcandı; kalan {o['kalan']:,.2f}."))
        if o["roas"] is not None and o["roas"] < roas_esik:
            uy.append(uyari("Orta", "ROAS düşük", kim, f"Harcanan her 1 liraya {o['roas']} lira gelir yazılmış (eşik {roas_esik}). Harcama {o['harcama']:,.2f}, gelir {o['gelir']:,.2f}."))
        if o["gelir"] is None and o["harcama"] > 0:
            uy.append(uyari("Bilgi", "Gelir yazılmamış", kim, f"{o['harcama']:,.2f} harcama var, gelir sütunu boş; getiri hesaplanamadı."))
        liste.append(o)
    liste.sort(key=lambda o: (katla(o["kanal"]), katla(o["kampanya"])))
    kanallar = {}
    for c in kampanyalar.values():
        k = kanallar.setdefault(katla(c["kanal"]), {"kanal": c["kanal"], "butce": Z, "satirlar": []})
        k["butce"] += c["butce"] or Z
        k["satirlar"] += c["satirlar"]
    aylar = defaultdict(list)
    for h in hareket:
        aylar[f"{h['tarih']:%Y-%m}" if h["tarih"] else "(tarih yok)"].append(h)
    return {"kampanyalar": liste, "kanallar": [{"kanal": k["kanal"], **_olc(k["satirlar"], k["butce"])} for _, k in sorted(kanallar.items())], "aylar": [{"ay": ay, **_olc(s, None)} for ay, s in sorted(aylar.items())],
            "toplam": _olc(hareket, sum((b["butce"] for b in butce), Z)), "beklenen": beklenen, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kampanyalar"
    basliklar = ["Bütçe", "Harcama", "Kalan", "Kullanım %", "Gelir", "ROAS", "ROI %", "Dönüşüm", "Dönüşüm Başı Maliyet"]
    deger = lambda o: [o["butce"], o["harcama"], o["kalan"], o["kullanim"], o["gelir"], o["roas"], o["roi"], o["donusum"], o["dbm"]]   # noqa: E731
    _baslik(ws, ["Kanal", "Kampanya"] + basliklar + ["Beklenen Kullanım %", "Fark (puan)", "Hız", "Dönem Sonu Tahmini"], (20, 28, 15, 15, 15, 10, 15, 8, 10, 10, 13, 11, 10, 11, 15))
    for o in r["kampanyalar"]:
        _satir(ws, [o["kanal"], o["kampanya"]] + deger(o) + [r["beklenen"] if o["fark"] is not None else None, o["fark"], o["hiz"], o["tahmin"]], para_s=(3, 4, 5, 6, 7, 8, 9, 11, 12, 13, 15),
               dolgu={5: KIRMIZI if o["kalan"] is not None and o["kalan"] < 0 else None, 14: HIZ_RENK.get(o["hiz"])})
    ws = wb.create_sheet("Kanallar")
    _baslik(ws, ["Kanal"] + basliklar, (22, 15, 15, 15, 10, 15, 8, 10, 10, 13))
    for o in r["kanallar"] + [{"kanal": "TOPLAM", **r["toplam"]}]:
        n = _satir(ws, [o["kanal"]] + deger(o), para_s=(2, 3, 4, 5, 6, 7, 8, 10), dolgu={4: KIRMIZI if o["kalan"] is not None and o["kalan"] < 0 else None})
    for h in ws[n]:
        h.font = Font(bold=True)
    ws = wb.create_sheet("Aylar")
    _baslik(ws, ["Ay", "Harcama", "Gelir", "ROAS", "Dönüşüm", "Dönüşüm Başı Maliyet"], (13, 15, 15, 8, 10, 13))
    for o in r["aylar"]:
        _satir(ws, [o["ay"], o["harcama"], o["gelir"], o["roas"], o["donusum"], o["dbm"]], para_s=(2, 3, 4, 6))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(butce_yol: Path, harcama_yol: Path, cikti: Path, donem=None, bugun=None, hiz: Decimal = Decimal(15), roas_esik: Decimal = Decimal(1)) -> dict:
    butce, uy1 = oku_butce(butce_yol)
    hareket, uy2 = oku_harcama(harcama_yol)
    r = hesapla(butce, hareket, donem, bugun, hiz, roas_esik)
    r["uyarilar"] = sirala(uy1 + uy2 + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Kanal ve kampanya bazında pazarlama bütçesini, harcamayı ve getiriyi izler.")
    p.add_argument("--butce", type=Path, help="Kanal, Kampanya, Bütçe")
    p.add_argument("--harcamalar", type=Path, help="Tarih, Kanal, Kampanya, Harcama, Gelir, Dönüşüm")
    p.add_argument("--donem-baslangic", help="Bütçe döneminin ilk günü (GG.AA.YYYY); bitişle birlikte verilirse harcama hızı hesaplanır")
    p.add_argument("--donem-bitis", help="Bütçe döneminin son günü (GG.AA.YYYY)")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--hiz", default="15", help="Kullanım beklenenden bu kadar puan saparsa uyarılır (varsayılan 15)")
    p.add_argument("--roas", default="1", help="ROAS bu değerin altındaysa uyarılır (varsayılan 1)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "pazarlama_butcesi_takibi.xlsx")
    a = p.parse_args(argv)
    if (a.butce is None) != (a.harcamalar is None):
        print("[X] --butce ve --harcamalar birlikte verilmelidir")
        return 2
    ornek = a.butce is None
    byol, hyol = (ORNEK / "butce.csv", ORNEK / "harcamalar.csv") if ornek else (a.butce, a.harcamalar)
    bas, bit = (a.donem_baslangic or "01.01.2026", a.donem_bitis or "31.12.2026") if ornek else (a.donem_baslangic, a.donem_bitis)
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    donem = (tarih(bas), tarih(bit)) if bas or bit else None
    hiz, roas = para(a.hiz), para(a.roas)
    if bugun is None or hiz is None or hiz <= 0 or roas is None or roas < 0 or (donem and (None in donem or donem[0] > donem[1])):
        print("[X] Tarihler GG.AA.YYYY; --donem-baslangic ve --donem-bitis birlikte, başlangıç bitişten önce; --hiz pozitif; --roas sıfır veya pozitif olmalı")
        return 2
    for y in (byol, hyol):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(byol, hyol, a.cikti, donem, bugun, hiz, roas)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {len(r['kampanyalar'])} kampanya · bütçe {t['butce']:,.2f} · harcama {t['harcama']:,.2f}" + (f" (%{t['kullanim']})" if t["kullanim"] is not None else "") + (f" · beklenen %{r['beklenen']}" if r["beklenen"] is not None else "")
          + (f" · ROAS {t['roas']}" if t["roas"] is not None else ""))
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
