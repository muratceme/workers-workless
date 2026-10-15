import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

O = main.ORNEK


def h(takip="T1", hasta="P1", gun=1, kod="A", adet=1, tani=("I20.9",), brans="Kardiyoloji", tutar=100, satir=2):
    return {"satir": satir, "takip": takip, "hasta": hasta, "tarih": date(2026, 10, gun), "kod": kod, "ad": kod, "adet": D(adet), "tanilar": list(tani),
            "brans": brans, "tutar": D(tutar), "ihlaller": [], "risk": D(0)}


def r(kod, tur, deger, kapsam="takip", aciklama=""):
    return {"kod": kod, "tur": tur, "deger": deger, "kapsam": kapsam, "aciklama": aciklama}


class SgkTesti(unittest.TestCase):
    def test_genel(self):
        k = [h(tani=()), h(gun=2, tani=("I209",)), h(gun=3, adet=0), h(gun=3, adet=0)]
        main.genel_kontrol(k)
        self.assertEqual(k[0]["ihlaller"], ["Tanı kodu yok"])
        self.assertEqual(k[1]["ihlaller"], ["ICD-10 biçimi geçersiz: 'I209'"])
        self.assertEqual(k[3]["ihlaller"], ["Adet 0 ya da negatif", "Mükerrer kalem (aynı takip, işlem ve tarih)"])

    def test_tani_ve_brans(self):
        k = [h(tani=("I25.1", "E11")), h(tani=("J06.9",)), h(brans="Dahiliye")]
        main.kural_kontrol(k, [r("A", "Tanı Gerekli", "I20 I25"), r("A", "Branş Şartı", "Kardiyoloji; Kalp ve Damar Cerrahisi")])
        self.assertEqual([len(x["ihlaller"]) for x in k], [0, 1, 1])
        self.assertIn("I20, I25", k[1]["ihlaller"][0])

    def test_adet_siniri(self):
        k = [h(adet=2, tutar=200), h(gun=1, adet=1, tutar=100, takip="T2"), h(gun=2, adet=1, tutar=100, takip="T3")]
        main.kural_kontrol(k, [r("A", "Adet Sınırı", "1")])
        self.assertEqual([x["risk"] for x in k], [D("100.00"), D(0), D(0)])
        k = [h(adet=1, takip="T1"), h(adet=1, takip="T2"), h(gun=2, adet=1, takip="T3")]
        main.kural_kontrol(k, [r("A", "Adet Sınırı", "1", "gun")])
        self.assertEqual([len(x["ihlaller"]) for x in k], [0, 1, 0])
        k = [h(gun=1), h(gun=20), h(gun=25)]
        main.kural_kontrol(k, [r("A", "Adet Sınırı", "2", "yil")])
        self.assertEqual([len(x["ihlaller"]) for x in k], [0, 0, 1])

    def test_tekrar_ve_birlikte(self):
        k = [h(gun=20, satir=3), h(gun=1, satir=2), h(gun=31, hasta="P2")]
        main.kural_kontrol(k, [r("A", "Tekrar Süresi", "30")])
        self.assertIn("aradan 19 gün", k[0]["ihlaller"][0])
        self.assertEqual((k[1]["ihlaller"], k[2]["ihlaller"]), ([], []))
        k = [h(kod="A"), h(kod="B"), h(kod="A", takip="T2", gun=1), h(kod="B", takip="T3", gun=1)]
        main.kural_kontrol(k, [r("A", "Birlikte Faturalanamaz", "B")])
        self.assertEqual([len(x["ihlaller"]) for x in k], [1, 0, 0, 0])
        k = [h(kod="A"), h(kod="A", takip="T2", gun=1), h(kod="B", takip="T3", gun=1)]
        main.kural_kontrol(k, [r("A", "Birlikte Faturalanamaz", "B", "gun")])
        self.assertEqual([len(x["ihlaller"]) for x in k], [1, 1, 0])

    def test_ornek_ve_rapor(self):
        with tempfile.TemporaryDirectory() as tmp:
            s = main.calistir(O / "hizmetler.csv", O / "kurallar.csv", Path(tmp) / "s.xlsx")
            hz = s["hizmetler"]
            self.assertEqual([len(k["ihlaller"]) for k in hz], [0, 1, 0, 1, 1, 0, 1, 2, 1])
            self.assertEqual(sum(k["risk"] for k in hz), D("2420.00"))
            ws = load_workbook(Path(tmp) / "s.xlsx")["Kurallar"]
            self.assertEqual([ws.cell(i, 6).value for i in range(2, 8)], [1, 1, 1, 1, 1, 1])

    def test_cli_ve_hatalar(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(main.main(["--kurallar", str(Path(tmp) / "yok.csv")]), 1)
            k = Path(tmp) / "k.csv"
            k.write_text("İşlem Kodu;Kural;Değer\nA;Yaş Sınırı;18\n", encoding="utf-8")
            self.assertEqual(main.main(["--kurallar", str(k)]), 1)
            k.write_text("İşlem Kodu;Kural;Değer\nA;Adet Sınırı;bir\n", encoding="utf-8")
            self.assertEqual(main.main(["--kurallar", str(k)]), 1)


if __name__ == "__main__":
    unittest.main()
