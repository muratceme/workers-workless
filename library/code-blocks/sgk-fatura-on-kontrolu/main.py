"""
SGK Fatura Ön Kontrolü — Workers / Workless kod bloğu
Anlaşmalı Kurumlar ve Faturalama › SGK Faturalama Uzmanı
Medula'ya gönderilecek hizmet kalemlerini kurumun tanımladığı SUT kurallarına göre ön kontrolden geçirir. Kurallar
kurallar.csv dosyasındadır; kodun içinde SUT maddesi yoktur.

  Kural türleri (kurallar.csv → İşlem Kodu, Kural, Değer, Kapsam, Açıklama):
    Tanı Gerekli          = kalemin tanı kodlarından biri Değer'deki ICD-10 öneklerinden biriyle başlamalı (ör. "I20 I25 R07")
    Adet Sınırı           = aynı işlemin toplam adedi ≤ Değer; Kapsam: Takip (varsayılan) / Gün (hasta + gün) / Yıl (hasta + takvim yılı)
    Tekrar Süresi         = aynı hasta için aynı işlem Değer günden önce tekrar faturalanamaz
    Branş Şartı           = kalemin branşı Değer'deki branşlardan biri olmalı ("Kardiyoloji; Kalp ve Damar Cerrahisi")
    Birlikte Faturalanamaz = Değer'deki işlem(ler)le aynı takipte (Kapsam: Gün → aynı hasta ve gün) faturalanamaz
  Genel kontroller: tanı kodu yok · ICD-10 biçimi geçersiz · mükerrer kalem (takip, işlem, tarih) · adet ≤ 0
  Riskli tutar: kalemin tamamı; adet sınırında aşan adet × birim tutar.
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
KIRMIZI, SARI, YESIL = "FDE2E1", "FFF4CE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS, SIFIR = Decimal("0.01"), Decimal(0)
ICD = re.compile(r"^[A-Z]\d{2}(\.\d{1,2})?$")
KURAL_TURLERI = {"tani gerekli": "Tanı Gerekli", "adet siniri": "Adet Sınırı", "tekrar suresi": "Tekrar Süresi", "brans sarti": "Branş Şartı",
                 "birlikte faturalanamaz": "Birlikte Faturalanamaz"}

HIZMET = {"takip": ("Takip No",), "hasta": ("Hasta No", "Hasta"), "tarih": ("İşlem Tarihi", "Tarih"), "kod": ("İşlem Kodu", "SUT Kodu", "Kod"),
          "ad": ("İşlem Adı",), "adet": ("Adet",), "tani": ("Tanı Kodları", "Tanı", "ICD"), "brans": ("Branş",), "tutar": ("Tutar",)}
KURAL = {"kod": ("İşlem Kodu", "SUT Kodu", "Kod"), "kural": ("Kural", "Kural Türü"), "deger": ("Değer",), "kapsam": ("Kapsam",), "aciklama": ("Açıklama", "Dayanak")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    if x is None or str(x).strip() == "":
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace(" ", "")
    if re.fullmatch(r"-?\d{1,3}(\.\d{3})+(,\d+)?", s) or "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tl(x) -> str:
    return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def liste(x) -> list[str]:
    return [k.strip() for k in re.split(r"[\s;,|]+", str(x or "")) if k.strip()]


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    if not s:
        raise ValueError(f"{yol.name}: tablo boş")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in sozluk.items()}
    eksik = [sozluk[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise ValueError(f"{yol.name}: gerekli sütun yok: {', '.join(eksik)}")
    return [{a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()} for r in s[1:]]


def hizmetleri_oku(yol: Path) -> list[dict]:
    out = []
    for i, r in enumerate(kayitlar(yol, HIZMET, ("takip", "hasta", "tarih", "kod")), 2):
        t = tarih(r["tarih"])
        if not t:
            raise ValueError(f"{yol.name} satır {i}: işlem tarihi okunamadı ('{r['tarih']}')")
        adet = para(r["adet"]) if str(r["adet"]).strip() else Decimal(1)
        out.append({"satir": i, "takip": str(r["takip"]).strip(), "hasta": str(r["hasta"]).strip(), "tarih": t, "kod": str(r["kod"]).strip().upper(),
                    "ad": str(r["ad"]), "adet": adet if adet is not None else SIFIR, "tanilar": [x.upper() for x in liste(r["tani"])],
                    "brans": str(r["brans"]).strip(), "tutar": para(r["tutar"]) or SIFIR, "ihlaller": [], "risk": SIFIR})
    return out


def kurallari_oku(yol: Path) -> list[dict]:
    out = []
    for i, r in enumerate(kayitlar(yol, KURAL, ("kod", "kural", "deger")), 2):
        tur = KURAL_TURLERI.get(katla(r["kural"]))
        if not tur:
            raise ValueError(f"{yol.name} satır {i}: kural türü tanınmadı ('{r['kural']}'); {', '.join(KURAL_TURLERI.values())}")
        deger = str(r["deger"]).strip()
        if tur in ("Adet Sınırı", "Tekrar Süresi") and para(deger) is None:
            raise ValueError(f"{yol.name} satır {i}: {tur} için sayısal değer gerekli")
        out.append({"kod": str(r["kod"]).strip().upper(), "tur": tur, "deger": deger, "kapsam": katla(r["kapsam"]) or "takip", "aciklama": str(r["aciklama"])})
    return out


def ihlal(k: dict, metin: str, risk: Decimal) -> None:
    k["ihlaller"].append(metin)
    k["risk"] = min(max(k["risk"], risk), k["tutar"]).quantize(KURUS)


def genel_kontrol(hiz: list[dict]) -> None:
    gorulen = set()
    for k in hiz:
        if not k["tanilar"]:
            ihlal(k, "Tanı kodu yok", k["tutar"])
        for t in k["tanilar"]:
            if not ICD.match(t):
                ihlal(k, f"ICD-10 biçimi geçersiz: '{t}'", k["tutar"])
        if k["adet"] <= 0:
            ihlal(k, "Adet 0 ya da negatif", k["tutar"])
        a = (k["takip"], k["kod"], k["tarih"])
        if a in gorulen:
            ihlal(k, "Mükerrer kalem (aynı takip, işlem ve tarih)", k["tutar"])
        gorulen.add(a)


def kural_kontrol(hiz: list[dict], kurallar: list[dict]) -> None:
    for r in kurallar:
        ilgili = [k for k in hiz if k["kod"] == r["kod"]]
        etiket = f"{r['tur']}" + (f" ({r['aciklama']})" if r["aciklama"] else "")
        if r["tur"] == "Tanı Gerekli":
            onek = [x.upper() for x in liste(r["deger"])]
            for k in ilgili:
                if k["tanilar"] and not any(t.startswith(o) for t in k["tanilar"] for o in onek):
                    ihlal(k, f"{etiket}: tanılar ({', '.join(k['tanilar'])}) {', '.join(onek)} ile uyumlu değil", k["tutar"])
                elif not k["tanilar"]:
                    ihlal(k, f"{etiket}: tanı yok", k["tutar"])
        elif r["tur"] == "Adet Sınırı":
            sinir = para(r["deger"])
            grup = defaultdict(lambda: SIFIR)
            for k in sorted(ilgili, key=lambda x: (x["tarih"], x["satir"])):
                anahtar = {"gun": (k["hasta"], k["tarih"]), "yil": (k["hasta"], k["tarih"].year)}.get(r["kapsam"], (k["takip"],))
                onceki = grup[anahtar]
                grup[anahtar] += k["adet"]
                if grup[anahtar] > sinir:
                    asan = min(k["adet"], grup[anahtar] - max(sinir, onceki))
                    birim = k["tutar"] / k["adet"] if k["adet"] else SIFIR
                    ihlal(k, f"{etiket}: sınır {sinir:g}/{r['kapsam']}, toplam {grup[anahtar]:g}", asan * birim)
        elif r["tur"] == "Tekrar Süresi":
            gun = int(para(r["deger"]))
            son = {}
            for k in sorted(ilgili, key=lambda x: (x["tarih"], x["satir"])):
                if k["hasta"] in son and (k["tarih"] - son[k["hasta"]]).days < gun:
                    ihlal(k, f"{etiket}: önceki işlem {son[k['hasta']]:%d.%m.%Y}, aradan {(k['tarih'] - son[k['hasta']]).days} gün geçmiş (en az {gun})",
                          k["tutar"])
                son[k["hasta"]] = k["tarih"]
        elif r["tur"] == "Branş Şartı":
            izinli = {katla(x) for x in re.split(r"[;|,]", r["deger"]) if x.strip()}
            for k in ilgili:
                if katla(k["brans"]) not in izinli:
                    ihlal(k, f"{etiket}: branş '{k['brans'] or '-'}'", k["tutar"])
        elif r["tur"] == "Birlikte Faturalanamaz":
            diger = {x.upper() for x in liste(r["deger"])}
            for k in ilgili:
                es = [x for x in hiz if x["kod"] in diger and ((x["hasta"], x["tarih"]) == (k["hasta"], k["tarih"]) if r["kapsam"] == "gun"
                                                                 else x["takip"] == k["takip"])]
                if es:
                    ihlal(k, f"{etiket}: aynı {'gün' if r['kapsam'] == 'gun' else 'takip'}te {', '.join(sorted({x['kod'] for x in es}))} var", k["tutar"])


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[1]:
        h.fill, h.font, h.alignment = BASLIK_DOLGU, BASLIK_YAZI, UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def rapor_yaz(cikti: Path, hiz: list[dict], kurallar: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kalemler"
    _baslik(ws, ["Takip No", "Hasta No", "Tarih", "İşlem Kodu", "İşlem", "Adet", "Tanılar", "Branş", "Tutar", "Riskli Tutar", "İhlaller"],
            (10, 9, 11, 10, 26, 6, 14, 16, 11, 11, 80))
    for k in hiz:
        ws.append([k["takip"], k["hasta"], k["tarih"], k["kod"], k["ad"], float(k["adet"]), ", ".join(k["tanilar"]), k["brans"], float(k["tutar"]),
                   float(k["risk"]), "\n".join(k["ihlaller"])])
        ws.cell(ws.max_row, 3).number_format = TF
        for c in (9, 10):
            ws.cell(ws.max_row, c).number_format = PF
        ws.cell(ws.max_row, 10).fill = PatternFill("solid", fgColor=KIRMIZI if k["ihlaller"] else YESIL)
        for h in ws[ws.max_row]:
            h.alignment = UST
    o = wb.create_sheet("Kurallar")
    _baslik(o, ["İşlem Kodu", "Kural", "Değer", "Kapsam", "Açıklama", "İhlal Sayısı"], (12, 22, 26, 10, 60, 12))
    for r in kurallar:
        n = sum(1 for k in hiz if k["kod"] == r["kod"] and any(i.startswith(r["tur"]) for i in k["ihlaller"]))
        o.append([r["kod"], r["tur"], r["deger"], r["kapsam"], r["aciklama"], n])
        if n:
            o.cell(o.max_row, 6).fill = PatternFill("solid", fgColor=KIRMIZI)
    top, risk = sum((k["tutar"] for k in hiz), SIFIR), sum((k["risk"] for k in hiz), SIFIR)
    o.append([])
    o.append([f"Toplam {tl(top)} TL · riskli {tl(risk)} TL · {sum(1 for k in hiz if k['ihlaller'])} / {len(hiz)} kalemde ihlal. Ön kontroldür; kuralları "
              "güncel Sağlık Uygulama Tebliği (SUT) ve eklerine göre kurumunuz tanımlar, Medula'nın kendi kontrolleri ayrıca uygulanır."])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(hizmet_yol: Path, kural_yol: Path, cikti: Path) -> dict:
    hiz = hizmetleri_oku(hizmet_yol)
    kurallar = kurallari_oku(kural_yol)
    genel_kontrol(hiz)
    kural_kontrol(hiz, kurallar)
    rapor_yaz(cikti, hiz, kurallar)
    return {"hizmetler": hiz, "kurallar": kurallar}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Medula'ya gönderilecek hizmet kalemlerini yapılandırılabilir SUT kurallarına göre ön kontrolden geçirir.")
    p.add_argument("--hizmetler", type=Path, default=ORNEK / "hizmetler.csv",
                   help="Takip No, Hasta No, İşlem Tarihi, İşlem Kodu, İşlem Adı, Adet, Tanı Kodları, Branş, Tutar")
    p.add_argument("--kurallar", type=Path, default=ORNEK / "kurallar.csv",
                   help="İşlem Kodu, Kural (Tanı Gerekli / Adet Sınırı / Tekrar Süresi / Branş Şartı / Birlikte Faturalanamaz), Değer, Kapsam, Açıklama")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "sgk_fatura_kontrol.xlsx")
    a = p.parse_args(argv)
    for yol in (a.hizmetler, a.kurallar):
        if not yol.exists():
            print(f"[X] Dosya bulunamadı: {yol}")
            return 1
    try:
        r = calistir(a.hizmetler, a.kurallar, a.cikti)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    hiz = r["hizmetler"]
    print(f"[OK] {len(hiz)} kalem · {len(r['kurallar'])} kural · ihlalli {sum(1 for k in hiz if k['ihlaller'])} kalem · "
          f"riskli {tl(sum((k['risk'] for k in hiz), SIFIR))} TL")
    for k in hiz:
        for i in k["ihlaller"]:
            print(f"[!] {k['takip']} {k['kod']}: {i}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
