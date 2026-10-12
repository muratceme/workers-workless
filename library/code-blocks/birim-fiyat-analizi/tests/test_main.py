import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def gd(tur, miktar, fiyat, poz="P1", rayic="", girdi="", kaynak="Analiz"):
    return {"poz": poz, "poz_tanim": "Tanım" if poz == "P1" else "", "olcu": "m³", "tur": tur, "rayic": rayic, "girdi": girdi or f"{tur}{miktar}", "birim": "", "miktar": D(miktar), "fiyat": D(fiyat), "kaynak": kaynak}


class AnalizTesti(unittest.TestCase):
    def test_poz_hesabi(self):
        l = [gd("Malzeme", "0.25", 4000), gd("Malzeme", "1.2", 600), gd("Malzeme", "0.15", 40), gd("İşçilik", 4, 250), gd("İşçilik", 1, 350), gd("Makine", "0.5", 800)]
        r = main.analiz(l)
        p = r["pozlar"][0]
        self.assertEqual((p["Malzeme"], p["İşçilik"], p["Makine"], p["Diğer"], p["toplam"], p["kar"], p["birim_fiyat"], p["girdi"], p["tanim"]), (D(1726), D(1350), D(400), D(0), D(3476), D(869), D(4345), 6, "Tanım"))
        self.assertEqual((p["pay"]["Malzeme"], p["pay"]["İşçilik"], p["pay"]["Makine"], l[0]["tutar"], l[0]["pay"]), (D("49.65"), D("38.84"), D("11.51"), D(1000), D("28.77")))
        self.assertEqual([main.analiz([gd("Malzeme", 1, 100)], D(k))["pozlar"][0]["birim_fiyat"] for k in (25, 0, "12.5")], [D(125), D(100), D("112.50")])
        self.assertEqual(main.analiz([gd("Malzeme", "0.333", "33.33")])["pozlar"][0]["kar"], D("2.78"))              # 11,10 × %25 = 2,775 → 2,78

    def test_birden_cok_poz_ve_uyarilar(self):
        r = main.analiz([gd("Malzeme", 1, 100), gd("Diğer", 2, 10, "P2", rayic="R1"), gd("İşçilik", 1, 0, "P2", rayic="R2", kaynak="Yok"), gd("Diğer", 1, 10, "P2", rayic="r1")])
        self.assertEqual([(p["poz"], p["toplam"], p["birim_fiyat"], p["eksik"]) for p in r["pozlar"]], [("P1", D(100), D(125), 0), ("P2", D(30), D("37.50"), 1)])
        self.assertEqual(([u["tur"] for u in r["uyarilar"]], r["eksik_poz"]), (["Aynı girdi iki kez"], 1))
        self.assertEqual([main.tur_coz(x) for x in ("MALZEME", "İşçilik", "işçi", "Makine", "Ekipman kirası", "Nakliye", "")], ["Malzeme", "İşçilik", "İşçilik", "Makine", "Makine", "Diğer", "Diğer"])
        self.assertEqual([main.miktar_coz(x) for x in ("0,25", "0.150", "1.250,5", "x")], [D("0.25"), D("0.150"), D("1250.5"), None])
        with self.assertRaises(ValueError):
            main.analiz([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "b.xlsx"
            r = main.calistir(ORNEK / "analiz_girdileri.csv", cikti, ORNEK / "rayicler.csv")
            self.assertEqual(([(p["poz"], p["olcu"], p["toplam"], p["kar"], p["birim_fiyat"]) for p in r["pozlar"]], r["uyarilar"]),
                             ([("POZ-01", "m³", D(3476), D(869), D(4345)), ("POZ-02", "m²", D(570), D("142.50"), D("712.50"))], []))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Poz Özeti", "Analiz Detayı", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            a, ry = Path(tmp) / "a.csv", Path(tmp) / "r.csv"
            ry.write_text("Rayiç No;Tanım;Birim;Fiyat\nR1;Çimento;ton;4.000,00\nR1;Çimento;ton;1\nR2;İşçi;saat;yok\n", encoding="utf-8")
            a.write_text("Poz No;Tür;Rayiç No;Girdi;Birim;Miktar;Birim Fiyat\nP1;Malzeme;R1;;kg;0.250;\nP1;İşçilik;R2;İşçi;saat;2;\nP1;Makine;;Vinç;saat;1;500\nP1;Malzeme;R1;;ton;1;3.900,00\nP1;Malzeme;R1;;ton;çok;\n",
                         encoding="utf-8")
            l, rayic, uy = main.oku(a, ry)
            self.assertEqual(([(g["girdi"], g["miktar"], g["fiyat"], g["kaynak"], g["birim"]) for g in l], rayic),
                             ([("Çimento", D("0.250"), D("4000.00"), "Rayiç", "kg"), ("İşçi", D(2), D(0), "Yok", "saat"), ("Vinç", D(1), D(500), "Analiz", "saat"), ("Çimento", D(1), D("3900.00"), "Analiz", "ton")],
                              {"r1": {"tanim": "Çimento", "birim": "ton", "fiyat": D("4000.00")}}))
            self.assertEqual(sorted(u["tur"] for u in uy), ["Analiz fiyatı rayiçten farklı", "Birim uyuşmuyor", "Miktar okunamadı", "Mükerrer rayiç", "Rayiç bulunamadı", "Rayiç fiyatı okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--analiz", str(a), "--cikti", c]
            self.assertEqual(main.main(arg + ["--rayicler", str(ry), "--kar-orani", "20"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(arg + ["--kar-orani", "eksi"]), 2)
            self.assertEqual(main.main(arg + ["--rayicler", str(Path(tmp) / "yok.csv")]), 1)
            a.write_text("Poz No;Miktar\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
