"""
Birim Fiyat Analizi — Workers / Workless kod bloğu
Teknik Ofis › Teknik Ofis Mühendisi

Malzeme, işçilik ve makine rayiçlerinden poz bazında birim fiyat analizi ve kâr-genel gider dahil fiyat oluşturur.
  Girdi tutarı       = miktar × birim fiyat                  (birim fiyat analiz satırında yazılıysa o, yoksa rayiç dosyasından rayiç numarasıyla alınır)
  Analiz toplamı     = pozun malzeme + işçilik + makine + diğer girdi tutarları
  Kâr ve genel gider = analiz toplamı × --kar-orani %        (varsayılan 25)
  Birim fiyat        = analiz toplamı + kâr ve genel gider
Miktarlarda virgül de nokta da ondalık ayırıcıdır (0,25 = 0.25). Fiyatlarda nokta binlik ayırıcı olabilir (4.000,00).
Rayiçler ve analiz miktarları sizin dosyanızdan gelir: kod hiçbir kurumun birim fiyat kitabını, analiz numarasını veya güncel rayicini bilmez.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


ANALIZ = {"poz": ("poz no", "poz", "analiz no", "is kalemi no"), "tanim": ("poz tanimi", "is kalemi", "isin adi", "poz adi"), "olcu": ("olcu birimi", "poz birimi"), "tur": ("tur", "girdi turu", "cins"),
          "rayic": ("rayic no", "rayic", "rayic kodu", "kod"), "girdi": ("girdi", "girdi tanimi", "aciklama", "tanim"), "birim": ("birim", "girdi birimi"), "miktar": ("miktar",),
          "fiyat": ("birim fiyat", "fiyat", "rayic fiyati")}
RAYIC = {"rayic": ANALIZ["rayic"], "tanim": ("tanim", "aciklama", "girdi", "rayic adi"), "birim": ("birim",), "fiyat": ("fiyat", "birim fiyat", "rayic fiyati")}
TURLER = ("Malzeme", "İşçilik", "Makine", "Diğer")
Z = Decimal(0)


def miktar_coz(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    s = s.replace(".", "").replace(",", ".") if "," in s and "." in s else s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tur_coz(x) -> str:
    k = katla(x)
    return "Malzeme" if k.startswith("malz") else "İşçilik" if k.startswith("isc") or "iscilik" in k else "Makine" if k.startswith("mak") or "ekipman" in k else "Diğer"


def oku(analiz_yol: Path, rayic_yol: Path | None = None) -> tuple[list[dict], dict, list[dict]]:
    uy, l, rayic = [], [], {}
    if rayic_yol:
        for r in kayitlar(rayic_yol, RAYIC, ("rayic", "fiyat")):
            k, f = metin(r["rayic"]), para(r["fiyat"])
            if not k:
                continue
            if f is None or f < 0:
                uy.append(uyari("Yüksek", "Rayiç fiyatı okunamadı", k, f"{rayic_yol.name} satır {r['_satir']}: '{metin(r['fiyat'])}'."))
                continue
            if katla(k) in rayic:
                uy.append(uyari("Yüksek", "Mükerrer rayiç", k, f"{rayic_yol.name}: aynı rayiç numarası birden fazla satırda; ilk satırdaki fiyat kullanıldı."))
                continue
            rayic[katla(k)] = {"tanim": metin(r["tanim"]), "birim": metin(r["birim"]), "fiyat": f}
    for r in kayitlar(analiz_yol, ANALIZ, ("poz", "miktar")):
        poz = metin(r["poz"])
        if not poz:
            continue
        kim = f"{poz} · {metin(r['girdi']) or metin(r['rayic']) or 'girdi'} (satır {r['_satir']})"
        m, f, kod = miktar_coz(r["miktar"]), para(r["fiyat"]), metin(r["rayic"])
        if m is None or m < 0:
            uy.append(uyari("Yüksek", "Miktar okunamadı", kim, f"'{metin(r['miktar'])}'. Girdi analize alınmadı; pozun birim fiyatı eksik çıkar."))
            continue
        ry = rayic.get(katla(kod)) if kod else None
        kaynak = "Analiz"
        if f is None:
            if metin(r["fiyat"]):
                uy.append(uyari("Yüksek", "Birim fiyat okunamadı", kim, f"'{metin(r['fiyat'])}'."))
            if ry:
                f, kaynak = ry["fiyat"], "Rayiç"
            else:
                uy.append(uyari("Yüksek", "Rayiç bulunamadı", kim, f"Birim fiyat yazılmamış ve rayiç '{kod or '(boş)'}' rayiç dosyasında yok. Girdi 0 fiyatla alındı; pozun birim fiyatı eksik çıkar."))
                f, kaynak = Z, "Yok"
        elif ry and ry["fiyat"] != f:
            uy.append(uyari("Orta", "Analiz fiyatı rayiçten farklı", kim, f"Analizde {f:f}, rayiç dosyasında {ry['fiyat']:f}. Analizdeki fiyat kullanıldı."))
        birim = metin(r["birim"]) or (ry["birim"] if ry else "")
        if ry and ry["birim"] and metin(r["birim"]) and katla(ry["birim"]) != katla(r["birim"]):
            uy.append(uyari("Orta", "Birim uyuşmuyor", kim, f"Analizde '{metin(r['birim'])}', rayiçte '{ry['birim']}'. Miktarın rayiç birimine göre yazıldığını kontrol edin."))
        l.append({"poz": poz, "poz_tanim": metin(r["tanim"]), "olcu": metin(r["olcu"]), "tur": tur_coz(r["tur"]), "rayic": kod, "girdi": metin(r["girdi"]) or (ry["tanim"] if ry else ""), "birim": birim,
                  "miktar": m, "fiyat": f, "kaynak": kaynak})
    return l, rayic, uy


def analiz(girdiler: list[dict], kar_orani: Decimal = Decimal(25)) -> dict:
    if not girdiler:
        raise ValueError("Analiz edilecek girdi yok.")
    pozlar, uy = {}, []
    for g in girdiler:
        g["tutar"] = yuvarla(g["miktar"] * g["fiyat"])
        p = pozlar.setdefault(katla(g["poz"]), {"poz": g["poz"], "tanim": "", "olcu": "", "girdi": 0, "eksik": 0, "gorulen": set()} | {t: Z for t in TURLER})
        p["tanim"], p["olcu"] = p["tanim"] or g["poz_tanim"], p["olcu"] or g["olcu"]
        p[g["tur"]] += g["tutar"]
        p["girdi"] += 1
        p["eksik"] += g["kaynak"] == "Yok"
        anahtar = katla(g["rayic"]) or katla(g["girdi"])
        if anahtar and anahtar in p["gorulen"]:
            uy.append(uyari("Orta", "Aynı girdi iki kez", f"{g['poz']} · {g['rayic'] or g['girdi']}", "Pozun analizinde aynı rayiç veya girdi birden fazla satırda; ikisi de toplandı. Mükerrer değilse dikkate almayın."))
        p["gorulen"].add(anahtar)
    for p in pozlar.values():
        p["toplam"] = sum((p[t] for t in TURLER), Z)
        p["kar"] = yuvarla(p["toplam"] * kar_orani / 100)
        p["birim_fiyat"] = p["toplam"] + p["kar"]
        p["pay"] = {t: yuvarla(p[t] / p["toplam"] * 100) if p["toplam"] else Z for t in TURLER}
    for g in girdiler:
        t = pozlar[katla(g["poz"])]["toplam"]
        g["pay"] = yuvarla(g["tutar"] / t * 100) if t else Z
    return {"pozlar": list(pozlar.values()), "girdiler": girdiler, "kar_orani": kar_orani, "uyarilar": uy, "eksik_poz": sum(p["eksik"] > 0 for p in pozlar.values())}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    ko = r["kar_orani"].normalize()
    ws = wb.active
    ws.title = "Poz Özeti"
    _baslik(ws, ["Poz No", "Poz Tanımı", "Ölçü Birimi", *TURLER, "Analiz Toplamı", f"Kâr ve Genel Gider (%{ko:f})", "Birim Fiyat", *[f"{t} Payı %" for t in TURLER[:3]], "Fiyatı Eksik Girdi"],
            (14, 44, 9, 14, 14, 14, 12, 15, 15, 15, 10, 10, 10, 10))
    for p in r["pozlar"]:
        _satir(ws, [p["poz"], p["tanim"], p["olcu"], *[p[t] for t in TURLER], p["toplam"], p["kar"], p["birim_fiyat"], *[p["pay"][t] for t in TURLER[:3]], p["eksik"] or None],
               para_s=tuple(range(4, 14)), dolgu={10: YESIL if not p["eksik"] else SARI, 14: KIRMIZI if p["eksik"] else None}, sar=(2,))
    ws.append([])
    ws.append([f"Birim fiyat = analiz toplamı × (1 + %{ko:f}). Rayiçler ve analiz miktarları dosyalarınızdan alınmıştır; kod resmi birim fiyat veya rayiç bilmez. Fiyatı eksik girdisi olan pozların birim fiyatı düşük çıkar."])
    ws = wb.create_sheet("Analiz Detayı")
    _baslik(ws, ["Poz No", "Tür", "Rayiç No", "Girdi", "Birim", "Miktar", "Birim Fiyat", "Tutar", "Poz İçindeki Payı %", "Fiyat Kaynağı"], (14, 10, 12, 40, 8, 10, 13, 14, 10, 11))
    for g in sorted(r["girdiler"], key=lambda g: (g["poz"], TURLER.index(g["tur"]))):
        n = _satir(ws, [g["poz"], g["tur"], g["rayic"], g["girdi"], g["birim"], g["miktar"], g["fiyat"], g["tutar"], g["pay"], g["kaynak"]], para_s=(7, 8, 9), dolgu={10: KIRMIZI if g["kaynak"] == "Yok" else None})
        ws.cell(n, 6).number_format = "0.000"
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(analiz_yol: Path, cikti: Path, rayic_yol: Path | None = None, kar_orani: Decimal = Decimal(25)) -> dict:
    girdiler, _, uy = oku(analiz_yol, rayic_yol)
    r = analiz(girdiler, kar_orani)
    uy += r["uyarilar"]
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Rayiçlerden poz bazında birim fiyat analizi yapar.")
    p.add_argument("--analiz", type=Path, default=ORNEK / "analiz_girdileri.csv", help="Poz No, Poz Tanımı, Ölçü Birimi, Tür (Malzeme / İşçilik / Makine), Rayiç No, Girdi, Birim, Miktar, Birim Fiyat")
    p.add_argument("--rayicler", type=Path, help="İsteğe bağlı: Rayiç No, Tanım, Birim, Fiyat")
    p.add_argument("--kar-orani", default="25", help="Kâr ve genel gider oranı %% (varsayılan 25)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "birim_fiyat_analizi.xlsx")
    a = p.parse_args(argv)
    rayic, ko = a.rayicler or (ORNEK / "rayicler.csv" if a.analiz == ORNEK / "analiz_girdileri.csv" else None), para(a.kar_orani)
    if ko is None or ko < 0:
        print("[X] --kar-orani sıfır veya pozitif sayı olmalı")
        return 2
    for y in (a.analiz, rayic):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.analiz, a.cikti, rayic, ko)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    for x in r["pozlar"]:
        print(f"[{'!' if x['eksik'] else 'OK'}] {x['poz']}: analiz toplamı {x['toplam']:,.2f} + kâr ve genel gider {x['kar']:,.2f} = birim fiyat {x['birim_fiyat']:,.2f}" + (f" / {x['olcu']}" if x["olcu"] else "")
              + (f" ({x['eksik']} girdinin fiyatı eksik)" if x["eksik"] else ""))
    print(f"[OK] {len(r['pozlar'])} poz · rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
