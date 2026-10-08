import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def calistir(cikti, **ek):
    return main.calistir(ORNEK / "provizyon_talepleri.csv", ORNEK / "policeler.csv", ORNEK / "plan_teminatlari.csv", cikti,
                         ORNEK / "genel_sartlar.csv", ORNEK / "onceki_odemeler.csv", **ek)


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "p.xlsx"
        cls.s = calistir(cls.cikti)
        cls.t = {t.no: t for t in cls.s["talepler"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_kararlar(self):
        self.assertEqual({k: t.karar for k, t in self.t.items()},
                         {"PV-01": "Onay", "PV-02": "Onay", "PV-03": "Ret", "PV-04": "Ret", "PV-05": "Kısmi Onay", "PV-06": "Ret", "PV-07": "İnceleme",
                          "PV-08": "Ret", "PV-09": "Ret", "PV-10": "İnceleme", "PV-11": "Ret", "PV-12": "İnceleme", "PV-13": "Ret"})

    def test_gerekceler(self):
        g = {k: t.gerekce[0] for k, t in self.t.items()}
        self.assertIn("365 gün", g["PV-03"])
        self.assertIn("205 gün", g["PV-03"])
        self.assertIn("ön mevcut durum istisnası: E11", g["PV-04"])
        self.assertIn("poliçe süresi dışında", g["PV-06"])
        self.assertIn("Prim durumu", g["PV-07"])
        self.assertIn("Genel şart istisnası: Z41.1", g["PV-08"])
        self.assertIn("H25", g["PV-09"])
        self.assertIn("128 gün", g["PV-09"])
        self.assertIn("18 gün sonra", g["PV-10"])
        self.assertIn("anlaşmalı değil", g["PV-11"])
        self.assertIn("ICD-10 tanı kodu yok", g["PV-12"])
        self.assertIn("'Diş' teminatı yok", g["PV-13"])

    def test_tutarlar(self):
        self.assertEqual((self.t["PV-02"].odenecek, self.t["PV-02"].katilim), (D("1920.00"), D("480.00")))
        pv5 = self.t["PV-05"]                    # 30.000 limit − 24.500 önceki = 5.500; 9.000 × 0,90 = 8.100
        self.assertEqual((pv5.kalan_once, pv5.odenecek, pv5.kalan_sonra), (D(5500), D("5500.00"), D(0)))
        self.assertIn("aşan 2.600,00 TL", pv5.gerekce[-1])
        self.assertEqual(self.t["PV-01"].odenecek, D("68000.00"))
        lim = {(x["sigortali"], x["teminat"]): x for x in self.s["limitler"]}
        self.assertEqual(lim[("S-01", "Yatarak Tedavi")]["kalan"], D(1432000))

    def test_excel(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Değerlendirme", "Limit Durumu", "Uyarılar"])
        ws = wb["Değerlendirme"]
        self.assertEqual(ws["R1"].value, "Uzman Kararı")
        self.assertIn("kısmen onaylanmıştır", ws["Q6"].value)


class KuralTesti(unittest.TestCase):
    def test_limit_doldu_ve_kronik_gun(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t / "o.csv").write_text("Sigortalı No;Teminat;Ödenen Tutar (TL)\nS-02;Ayakta Tedavi;15000\n", encoding="utf-8")
            s = main.calistir(ORNEK / "provizyon_talepleri.csv", ORNEK / "policeler.csv", ORNEK / "plan_teminatlari.csv", t / "x.xlsx",
                              ORNEK / "genel_sartlar.csv", t / "o.csv", kronik_gun=10)
            r = {x.no: x for x in s["talepler"]}
            self.assertEqual(r["PV-02"].karar, "Ret")
            self.assertIn("limiti", r["PV-02"].gerekce[0])
            self.assertEqual(r["PV-10"].karar, "Onay")          # 18 gün > 10

    def test_icd(self):
        k = main.Kural("E11", "Kronik", None, "")
        self.assertTrue(k.uyar("E11.65"))
        self.assertFalse(k.uyar("E10.1"))
        self.assertEqual(main.Police("p", "s", "", "", None, None, None, "", "E11 ve I10.0 istisna").istisna_kodlari, ["E11", "I10.0"])

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--policeler", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
