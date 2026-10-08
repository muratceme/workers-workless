"""
Provizyon Talebi Değerlendirme — Workers / Workless kod bloğu
Sigortacılık › Sağlık Sigortaları › Provizyon Uzmanı

Anlaşmalı kurumlardan gelen provizyon taleplerini poliçe, plan teminatları ve genel şart kurallarına göre ön
değerlendirir. Talepler tarih sırasıyla işlenir; onaylananlar kalan limitten düşülür.
  Karar sırası (ilk uyan karar verir):
  1. Anlaşmasız kurum → Ret (provizyon verilmez; geri ödeme yoluyla başvurulabilir)
  2. Poliçe bulunamadı / talep tarihi poliçe süresi dışında → Ret
  3. Teminat planda yok → Ret
  4. Genel şart istisnası (ICD-10 öneki) → Ret
  5. Beyan edilen ön mevcut durum istisnası (poliçedeki ICD-10 kodları) → Ret
  6. Teminat veya tanıya özel bekleme süresi dolmamış (ilk giriş tarihinden itibaren) → Ret
  7. ICD-10 kodu yok, prim gecikmede, poliçe girişine yakın (varsayılan 90 gün) kronik tanı → İnceleme
  8. Ödenecek = talep × (1 − katılım payı), kalan yıllık limitle sınırlı → Onay / Kısmi Onay / Limit doldu
Rapor: değerlendirme, limit durumu, uyarılar, kurum için kısa provizyon cevabı. İnternete bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: 13 talep, 7 poliçe
    python main.py --talepler talepler.xlsx --policeler policeler.xlsx --plan plan_teminatlari.csv --kurallar genel_sartlar.csv
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
ICD_RE = re.compile(r"\b([A-TV-Z]\d{2}(?:\.\d{1,2})?)\b")
EVET = ("evet", "e", "x", "var", "yes", "1", "anlasmali")
K2 = Decimal("0.01")

TALEP_SUTUNLARI = {"no": ("talep no", "provizyon no", "no"), "tarih": ("talep tarihi", "tarih", "islem tarihi"), "kurum": ("kurum", "saglik kurumu"),
                   "anlasmali": ("anlasmali", "anlasmali kurum"), "sigortali": ("sigortali no", "sigortali"), "police": ("police no", "police"),
                   "teminat": ("teminat",), "icd": ("icd 10", "icd", "tani kodu"), "tani": ("tani", "teshis"), "islem": ("islem", "aciklama"),
                   "tutar": ("talep tutari tl", "talep tutari", "tutar")}
POLICE_SUTUNLARI = {"police": ("police no", "police"), "sigortali": ("sigortali no", "sigortali"), "ad": ("ad soyad", "sigortali adi"),
                    "plan": ("plan", "urun"), "baslangic": ("baslangic", "baslangic tarihi"), "bitis": ("bitis", "bitis tarihi"),
                    "ilk_giris": ("ilk giris tarihi", "ilk giris", "sigortalilik baslangici"), "prim": ("prim durumu", "prim"),
                    "beyan": ("beyan on mevcut durum istisnasi", "on mevcut durum", "beyan", "ozel istisna")}
PLAN_SUTUNLARI = {"plan": ("plan",), "teminat": ("teminat",), "limit": ("yillik limit tl", "yillik limit", "limit"),
                  "katilim": ("katilim payi", "katilim payi yuzde", "katilim"), "bekleme": ("bekleme suresi gun", "bekleme suresi", "bekleme")}
KURAL_SUTUNLARI = {"anahtar": ("icd 10 anahtar", "icd 10", "icd", "anahtar"), "kapsam": ("kapsam", "tur"), "bekleme": ("bekleme suresi gun", "bekleme suresi"),
                   "aciklama": ("aciklama",)}
ODEME_SUTUNLARI = {"sigortali": ("sigortali no", "sigortali"), "teminat": ("teminat",), "tutar": ("odenen tutar tl", "odenen tutar", "tutar")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("₺", "").replace(" ", "").replace("%", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tl(x: Decimal | None) -> str:
    return "—" if x is None else f"{x.quantize(K2, ROUND_HALF_UP):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def icd_norm(s: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(s or "").upper())


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Police:
    no: str
    sigortali: str
    ad: str
    plan: str
    baslangic: date | None
    bitis: date | None
    ilk_giris: date | None
    prim: str
    beyan: str

    @property
    def istisna_kodlari(self) -> list[str]:
        return ICD_RE.findall(self.beyan.upper())


@dataclass
class Teminat:
    plan: str
    ad: str
    limit: Decimal | None
    katilim: Decimal
    bekleme: int | None


@dataclass
class Kural:
    anahtar: str
    kapsam: str           # İstisna / Bekleme / Kronik
    bekleme: int | None
    aciklama: str

    def uyar(self, icd: str) -> bool:
        return bool(icd) and icd_norm(icd).startswith(icd_norm(self.anahtar))


@dataclass
class Talep:
    satir: int
    no: str
    tarih: date | None
    kurum: str
    anlasmali: bool
    sigortali: str
    police: str
    teminat: str
    icd: str
    tani: str
    islem: str
    tutar: Decimal | None
    karar: str = ""
    gerekce: list = field(default_factory=list)
    katilim: Decimal = Decimal(0)
    odenecek: Decimal = Decimal(0)
    kalan_once: Decimal | None = None
    kalan_sonra: Decimal | None = None


def oku_policeler(yol: Path) -> dict[str, Police]:
    return {metin(r["police"]): Police(metin(r["police"]), metin(r.get("sigortali")), metin(r.get("ad")), metin(r.get("plan")), tarih(r.get("baslangic")),
                                       tarih(r.get("bitis")), tarih(r.get("ilk_giris")) or tarih(r.get("baslangic")), metin(r.get("prim")), metin(r.get("beyan")))
            for r in kayitlar(yol, POLICE_SUTUNLARI, ("police", "plan")) if r.get("police")}


def oku_plan(yol: Path) -> dict[tuple[str, str], Teminat]:
    sonuc = {}
    for r in kayitlar(yol, PLAN_SUTUNLARI, ("plan", "teminat")):
        if r.get("plan") and r.get("teminat"):
            k = para(r.get("katilim")) or Decimal(0)
            b = para(r.get("bekleme"))
            sonuc[(katla(r["plan"]), katla(r["teminat"]))] = Teminat(metin(r["plan"]), metin(r["teminat"]), para(r.get("limit")),
                                                                    k / 100 if k > 1 else k, int(b) if b else None)
    return sonuc


def oku_kurallar(yol: Path | None) -> list[Kural]:
    if yol is None:
        return []
    sonuc = []
    for r in kayitlar(yol, KURAL_SUTUNLARI, ("anahtar", "kapsam")):
        if r.get("anahtar"):
            k = katla(r.get("kapsam"))
            kapsam = "İstisna" if k.startswith("istisna") or k.startswith("kapsam disi") else "Bekleme" if k.startswith("bekleme") else "Kronik"
            b = para(r.get("bekleme"))
            sonuc.append(Kural(metin(r["anahtar"]), kapsam, int(b) if b else None, metin(r.get("aciklama"))))
    return sonuc


def oku_odemeler(yol: Path | None) -> dict[tuple[str, str], Decimal]:
    sonuc = defaultdict(Decimal)
    if yol is None:
        return sonuc
    for r in kayitlar(yol, ODEME_SUTUNLARI, ("sigortali", "teminat", "tutar")):
        if r.get("sigortali"):
            sonuc[(metin(r["sigortali"]), katla(r["teminat"]))] += para(r.get("tutar")) or Decimal(0)
    return sonuc


def oku_talepler(yol: Path) -> list[Talep]:
    return [Talep(r["_satir"], metin(r.get("no")) or f"Satır {r['_satir']}", tarih(r.get("tarih")), metin(r.get("kurum")),
                  katla(r.get("anlasmali") or "evet") in EVET, metin(r.get("sigortali")), metin(r.get("police")), metin(r.get("teminat")),
                  metin(r.get("icd")).upper(), metin(r.get("tani")), metin(r.get("islem")), para(r.get("tutar")))
            for r in kayitlar(yol, TALEP_SUTUNLARI, ("teminat", "tutar")) if r.get("no") or r.get("tutar")]


# ----------------------------------------------------------------------------
# Değerlendirme
# ----------------------------------------------------------------------------

def degerlendir(talepler: list[Talep], policeler: dict[str, Police], plan: dict, kurallar: list[Kural], odemeler: dict,
                kronik_gun: int = 90) -> dict:
    kullanilan = defaultdict(Decimal, odemeler)
    uyarilar = []
    sig_police = defaultdict(list)
    for p in policeler.values():
        sig_police[p.sigortali].append(p)
    for t in sorted(talepler, key=lambda t: (t.tarih or date.max, t.no)):
        def karar(k, neden):
            t.karar = k
            t.gerekce.append(neden)

        p = policeler.get(t.police)
        if p is None and t.sigortali in sig_police and t.tarih:
            p = next((x for x in sig_police[t.sigortali] if x.baslangic and x.bitis and x.baslangic <= t.tarih <= x.bitis), None)
        if p and t.sigortali and p.sigortali and p.sigortali != t.sigortali:
            uyarilar.append({"onem": "Yüksek", "talep": t.no, "aciklama": f"Sigortalı no ({t.sigortali}) poliçedeki sigortalıyla ({p.sigortali}) uyuşmuyor"})
        if t.tutar is None or t.tarih is None:
            karar("İnceleme", "Talep tutarı veya tarihi okunamadı")
            continue
        if not t.anlasmali:
            karar("Ret", f"{t.kurum or 'Kurum'} anlaşmalı değil; provizyon verilemez. Sigortalı faturayla geri ödeme başvurusu yapabilir.")
            continue
        if p is None:
            karar("Ret", f"Poliçe bulunamadı ({t.police or t.sigortali})")
            continue
        if not (p.baslangic and p.bitis and p.baslangic <= t.tarih <= p.bitis):
            karar("Ret", f"Talep tarihi {t.tarih:%d.%m.%Y} poliçe süresi dışında ({p.baslangic:%d.%m.%Y} – {p.bitis:%d.%m.%Y})")
            continue
        tem = plan.get((katla(p.plan), katla(t.teminat)))
        if tem is None:
            karar("Ret", f"'{p.plan}' planında '{t.teminat}' teminatı yok")
            continue
        istisna = next((k for k in kurallar if k.kapsam == "İstisna" and k.uyar(t.icd)), None)
        if istisna:
            karar("Ret", f"Genel şart istisnası: {t.icd} — {istisna.aciklama}")
            continue
        onmevcut = next((k for k in p.istisna_kodlari if t.icd and icd_norm(t.icd).startswith(icd_norm(k))), None)
        if onmevcut:
            karar("Ret", f"Poliçede ön mevcut durum istisnası: {onmevcut} ({p.beyan})")
            continue
        gecen = (t.tarih - p.ilk_giris).days if p.ilk_giris else None
        bekleme = [(tem.bekleme, f"{t.teminat} teminatı")] if tem.bekleme else []
        bekleme += [(k.bekleme, f"{k.anahtar} — {k.aciklama}") for k in kurallar if k.kapsam == "Bekleme" and k.bekleme and k.uyar(t.icd)]
        dolmayan = [(g, ad) for g, ad in bekleme if gecen is not None and gecen < g]
        if dolmayan:
            g, ad = max(dolmayan)
            karar("Ret", f"Bekleme süresi dolmamış: {ad} için {g} gün, ilk girişten ({p.ilk_giris:%d.%m.%Y}) bu yana {gecen} gün")
            continue
        inceleme = []
        if not t.icd:
            inceleme.append("ICD-10 tanı kodu yok; kurumdan tanı kodu ve epikriz isteyin")
        elif not ICD_RE.fullmatch(t.icd):
            inceleme.append(f"ICD-10 kodu '{t.icd}' biçimsel olarak geçersiz")
        if katla(p.prim).startswith(("gecik", "odenmedi", "borc")):
            inceleme.append(f"Prim durumu: {p.prim} — ödeme durumu ve poliçe şartları kontrol edilmeli")
        kronik = next((k for k in kurallar if k.kapsam == "Kronik" and k.uyar(t.icd)), None)
        if kronik and gecen is not None and gecen <= kronik_gun:
            inceleme.append(f"Kronik tanı ({t.icd} {kronik.aciklama}) ilk girişten {gecen} gün sonra: ön mevcut durum araştırması (önceki raporlar, "
                            "ilaç kullanımı) gerekebilir")
        # Tutar
        anahtar = (p.sigortali or t.sigortali, katla(t.teminat))
        brut = t.tutar * (1 - tem.katilim)
        kalan = (tem.limit - kullanilan[anahtar]) if tem.limit is not None else None
        t.kalan_once = kalan
        t.odenecek = max(min(brut, kalan) if kalan is not None else brut, Decimal(0)).quantize(K2, ROUND_HALF_UP)
        t.katilim = (t.tutar * tem.katilim).quantize(K2, ROUND_HALF_UP)
        if inceleme:
            t.karar = "İnceleme"
            t.gerekce += inceleme
            t.gerekce.append(f"Kurallara göre ödenebilir tutar {tl(t.odenecek)} TL (karar uzman / hekim incelemesinden sonra)")
            t.kalan_sonra = kalan
            continue
        if kalan is not None and kalan <= 0:
            karar("Ret", f"{t.teminat} yıllık limiti ({tl(tem.limit)} TL) dolmuş")
            t.odenecek = Decimal(0)
            continue
        t.karar = "Kısmi Onay" if kalan is not None and brut > kalan else "Onay"
        t.gerekce.append(f"{t.teminat} teminatı kapsamında" + (f"; kalan yıllık limit {tl(kalan)} TL" if kalan is not None else ""))
        if tem.katilim:
            t.gerekce.append(f"Katılım payı %{(tem.katilim * 100).normalize():f}: {tl(t.katilim)} TL sigortalıya ait")
        if t.karar == "Kısmi Onay":
            t.gerekce.append(f"Kalan yıllık limit {tl(kalan)} TL; aşan {tl(brut - kalan)} TL sigortalıya ait")
        kullanilan[anahtar] += t.odenecek
        t.kalan_sonra = kalan - t.odenecek if kalan is not None else None
        if t.tutar >= Decimal("100000"):
            uyarilar.append({"onem": "Bilgi", "talep": t.no, "aciklama": f"Yüksek tutarlı talep ({tl(t.tutar)} TL): kurum fiyat listesi ve tıbbi "
                                                                          "gereklilik kontrolü önerilir"})
    limitler = []
    for (sig, tem_k), harcanan in sorted(kullanilan.items()):
        p = next((x for x in policeler.values() if x.sigortali == sig), None)
        tem = plan.get((katla(p.plan), tem_k)) if p else None
        if tem:
            limitler.append({"sigortali": sig, "ad": p.ad, "plan": p.plan, "teminat": tem.ad, "limit": tem.limit, "kullanilan": harcanan,
                             "kalan": tem.limit - harcanan if tem.limit is not None else None})
    return {"talepler": sorted(talepler, key=lambda t: (t.tarih or date.max, t.no)), "limitler": limitler, "uyarilar": uyarilar}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Onay": "E3F4E1", "Kısmi Onay": "FFF4CE", "İnceleme": "FFF4CE", "Ret": "FDE2E1", "Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KONTROL = PatternFill("solid", fgColor="FFF4CE")
UST = Alignment(vertical="top", wrap_text=True)
PF = "#,##0.00"


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def cevap_metni(t: Talep) -> str:
    bas = f"{t.no} · {t.tani or t.icd} · {t.islem}".strip(" ·")
    if t.karar in ("Onay", "Kısmi Onay"):
        return (f"{bas}: provizyon {'kısmen ' if t.karar == 'Kısmi Onay' else ''}onaylanmıştır. Şirket payı {tl(t.odenecek)} TL"
                + (f", sigortalı payı {tl(t.tutar - t.odenecek)} TL" if t.tutar and t.tutar > t.odenecek else "") + ".")
    if t.karar == "Ret":
        return f"{bas}: provizyon verilememiştir. Gerekçe: {t.gerekce[0]}"
    return f"{bas}: talep incelemeye alınmıştır; ek bilgi / belge istenebilir."


def rapor_yaz(cikti: Path, s: dict) -> None:
    wb = Workbook()
    dg = wb.active
    dg.title = "Değerlendirme"
    _baslik(dg, ["Talep No", "Tarih", "Kurum", "Sigortalı", "Poliçe", "Teminat", "ICD-10", "Tanı", "İşlem", "Talep Tutarı", "Katılım Payı", "Ödenecek",
                 "Kalan Limit (önce)", "Kalan Limit (sonra)", "Ön Karar", "Gerekçe", "Kurum Cevabı (taslak)", "Uzman Kararı"],
            (9, 11, 20, 9, 9, 15, 8, 22, 26, 12, 11, 12, 13, 13, 11, 60, 60, 16))
    for t in s["talepler"]:
        dg.append([t.no, t.tarih, t.kurum, t.sigortali, t.police, t.teminat, t.icd, t.tani, t.islem, t.tutar and float(t.tutar), float(t.katilim),
                   float(t.odenecek), t.kalan_once and float(t.kalan_once), t.kalan_sonra and float(t.kalan_sonra), t.karar, "\n".join(t.gerekce),
                   cevap_metni(t), ""])
        r = dg.max_row
        dg.cell(r, 2).number_format = "DD.MM.YYYY"
        for j in (10, 11, 12, 13, 14):
            dg.cell(r, j).number_format = PF
        dg.cell(r, 15).fill = PatternFill("solid", fgColor=RENK[t.karar])
        dg.cell(r, 18).fill = KONTROL
        for j in (16, 17):
            dg.cell(r, j).alignment = UST
    dg.auto_filter.ref = f"A1:R{dg.max_row}"
    dg.append([])
    ozet = defaultdict(lambda: [0, Decimal(0)])
    for t in s["talepler"]:
        ozet[t.karar][0] += 1
        ozet[t.karar][1] += t.odenecek if t.karar in ("Onay", "Kısmi Onay") else 0
    dg.append(["Özet"] + [f"{k}: {v[0]} talep" + (f", {tl(v[1])} TL" if v[1] else "") for k, v in ozet.items()])

    lm = wb.create_sheet("Limit Durumu")
    _baslik(lm, ["Sigortalı", "Ad Soyad", "Plan", "Teminat", "Yıllık Limit", "Kullanılan (önceki + bu çalışma)", "Kalan", "Kullanım %"],
            (9, 18, 10, 16, 13, 16, 13, 10))
    for x in s["limitler"]:
        lm.append([x["sigortali"], x["ad"], x["plan"], x["teminat"], x["limit"] and float(x["limit"]), float(x["kullanilan"]), x["kalan"] and float(x["kalan"]),
                   float(x["kullanilan"] / x["limit"]) if x["limit"] else None])
        for j in (5, 6, 7):
            lm.cell(lm.max_row, j).number_format = PF
        lm.cell(lm.max_row, 8).number_format = "0%"

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Talep", "Açıklama", "İnceleme"], (9, 9, 90, 24))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["talep"], u["aciklama"], ""])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).fill = KONTROL
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(talep_yolu: Path, police_yolu: Path, plan_yolu: Path, cikti: Path, kural_yolu: Path | None = None, odeme_yolu: Path | None = None,
             kronik_gun: int = 90) -> dict:
    s = degerlendir(oku_talepler(talep_yolu), oku_policeler(police_yolu), oku_plan(plan_yolu), oku_kurallar(kural_yolu), oku_odemeler(odeme_yolu),
                    kronik_gun)
    rapor_yaz(cikti, s)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Provizyon taleplerini poliçe teminatı, limit, bekleme süresi ve istisnalara göre ön değerlendirir.")
    p.add_argument("--talepler", type=Path, default=ORNEK / "provizyon_talepleri.csv")
    p.add_argument("--policeler", type=Path, default=ORNEK / "policeler.csv")
    p.add_argument("--plan", type=Path, default=ORNEK / "plan_teminatlari.csv", help="Plan, Teminat, Yıllık Limit, Katılım Payı %%, Bekleme Süresi")
    p.add_argument("--kurallar", type=Path, default=ORNEK / "genel_sartlar.csv", help="ICD-10 öneki, Kapsam (İstisna / Bekleme / Kronik), Bekleme Süresi")
    p.add_argument("--odemeler", type=Path, help="İsteğe bağlı: Sigortalı No, Teminat, Ödenen Tutar (poliçe yılı içinde)")
    p.add_argument("--kronik-gun", type=int, default=90, help="İlk girişten bu kadar gün içindeki kronik tanılar incelemeye alınır (varsayılan 90)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "provizyon_degerlendirme.xlsx")
    a = p.parse_args(argv)
    if a.talepler == ORNEK / "provizyon_talepleri.csv" and a.odemeler is None:
        a.odemeler = ORNEK / "onceki_odemeler.csv"
    for y in (a.talepler, a.policeler, a.plan, a.kurallar, a.odemeler):
        if y is not None and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        s = calistir(a.talepler, a.policeler, a.plan, a.cikti, a.kurallar, a.odemeler, a.kronik_gun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    for t in s["talepler"]:
        isaret = {"Onay": "OK", "Kısmi Onay": "OK", "İnceleme": "!", "Ret": "X"}[t.karar]
        print(f"[{isaret}] {t.no} {t.karar:10} {tl(t.odenecek):>12} TL · {t.gerekce[0] if t.gerekce else ''}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
