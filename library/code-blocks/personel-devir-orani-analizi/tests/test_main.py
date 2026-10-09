import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BAS, BIT = date(2026, 1, 1), date(2026, 9, 30)


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "d.xlsx"
        cls.s = main.calistir(ORNEK / "personel.csv", cls.cikti, BAS, BIT)
        cls.g = cls.s["genel"]

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_genel(self):
        g = self.g
        self.assertEqual((g["bas"], g["son"], g["giren"], g["ayrilan"]), (51, 43, 11, 19))
        self.assertAlmostEqual(g["ort"], 48.6)                       # 10 ölçümün ortalaması (bağımsız hesapla doğrulandı)
        self.assertEqual((g["Gönüllü"], g["Gönülsüz"], g["Diğer"]), (14, 3, 2))
        self.assertAlmostEqual(g["yillik"], 19 / 48.6 * 365 / 273)
        self.assertAlmostEqual(g["tutma"], 36 / 51)

    def test_kirilimlar(self):
        y = {x["ad"]: x for x in self.s["yonetici"]}
        self.assertEqual((y["Yönetici B"]["ayrilan"], y["Yönetici B"]["Gönüllü"]), (8, 6))
        k = {x["ad"]: x for x in self.s["kidem"]}
        self.assertEqual((k["0–3 ay"]["ayrilan"], k["3–12 ay"]["ayrilan"], k["5 yıl +"]["ayrilan"]), (4, 1, 14))
        self.assertEqual([x["ayrilan"] for x in self.s["aylik"]], [2, 3, 1, 1, 1, 4, 1, 4, 2])

    def test_erken_ve_sinyaller(self):
        self.assertEqual((len(self.s["kohort"]), sorted(k.sicil for k in self.s["erken"])), (10, ["3050", "3051", "3056", "3058"]))
        self.assertEqual(len(self.s["ilk_yil"]), 5)
        u = {(x["tur"], x["kim"]) for x in self.s["uyarilar"]}
        self.assertTrue({("Erken ayrılma", "Yönetici B"), ("Erken ayrılma", "Yönetici D"), ("Yüksek gönüllü devir (departman)", "Depo"),
                         ("Yüksek gönüllü devir (yönetici)", "Yönetici B")} <= u)
        self.assertNotIn(("Yüksek gönüllü devir (yönetici)", "Yönetici C"), u)

    def test_excel(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Özet", "Kırılımlar", "Aylık Trend", "Erken Ayrılma", "Ayrılanlar", "Uyarılar"])
        self.assertEqual(wb["Ayrılanlar"].max_row, 20)
        self.assertEqual(wb["Özet"]["B7"].value, 19)


class KuralTesti(unittest.TestCase):
    def test_tur_bul(self):
        self.assertEqual(main.tur_bul("İstifa - daha yüksek ücret"), "Gönüllü")
        self.assertEqual(main.tur_bul("Deneme süresinde işveren feshi"), "Gönülsüz")
        self.assertEqual(main.tur_bul("İşçi tarafından haklı nedenle fesih"), "Gönüllü")
        self.assertEqual(main.tur_bul("Belirli süreli sözleşmenin sona ermesi"), "Diğer")
        self.assertEqual(main.tur_bul("Emeklilik"), "Diğer")
        self.assertEqual(main.tur_bul("?"), "Belirsiz")
        self.assertEqual(main.tur_bul("İstifa", "Gönülsüz"), "Gönülsüz")

    def test_ay_sonlari_ve_kucuk_grup(self):
        self.assertEqual(main.ay_sonlari(date(2026, 1, 15), date(2026, 3, 10)), [date(2026, 1, 31), date(2026, 2, 28), date(2026, 3, 10)])
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t / "p.csv").write_text("Sicil No;Departman;İşe Giriş;Ayrılış Tarihi;Ayrılış Nedeni\n1;X;01.01.2020;15.03.2026;Bilinmiyor\n"
                                     "2;X;01.01.2020;;\n3;Y;01.05.2026;01.04.2026;İstifa\n", encoding="utf-8")
            s = main.calistir(t / "p.csv", t / "o.xlsx", date(2026, 1, 1), date(2026, 6, 30))
            turler = {u["tur"] for u in s["uyarilar"]}
            self.assertTrue({"Ayrılış türü belirsiz", "Küçük grup", "Tarih hatası"} <= turler)
            self.assertEqual(s["genel"]["ayrilan"], 1)

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--donem", "30.09.2026", "01.01.2026"]), 2)
            self.assertEqual(main.main(["--personel", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
