"""
Personel Devir Oranı Analizi — Workers / Workless kod bloğu
İnsan Kaynakları › İnsan Kaynakları Müdürü

Giriş-çıkış verisinden personel devir (turnover) oranlarını ve erken ayrılma sinyallerini çıkarır:
  - Ortalama personel: dönem başı ve her ay sonu aktif personel sayılarının ortalaması.
    Ayrılış tarihi son çalışma günüdür; o günün sonunda çalışan aktif sayılmaz.
  - Devir oranı = dönemde ayrılan ÷ ortalama personel; yıllıklandırılmış = × 365 ÷ dönem günü.
  - Ayrılış türü: "Ayrılış Türü" sütunundan ya da nedenden (istifa → Gönüllü; işveren feshi, performans, disiplin →
    Gönülsüz; emeklilik, askerlik, ölüm, belirli süreli sözleşme sonu → Diğer). Tanınmayan neden "Belirsiz".
  - Kırılımlar: departman, yönetici, kıdem bandı (payda o banttaki ortalama personel), aylık trend.
  - Erken ayrılma: ayrılanlardan ilk yılını doldurmayanlar; dönemde işe alınıp 90 günü gözlenebilenlerden 90 gün
    içinde ayrılanlar (yönetici bazında).
  - Sinyal: gönüllü devir oranı şirket ortalamasının 1,5 katını aşan ve en az 2 gönüllü ayrılışı olan departman /
    yönetici; 2 ve üzeri erken ayrılış görülen yönetici.
Rapor: özet, kırılımlar, aylık trend, erken ayrılma, ayrılan listesi, uyarılar. İnternete bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: 63 kişi, 01.01–30.09.2026
    python main.py --personel personel.xlsx --donem 01.01.2026 31.12.2026
    python main.py --personel personel.xlsx --donem 01.01.2026 30.06.2026 --erken-gun 60 --kat 2
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")

SUTUNLAR = {"sicil": ("sicil no", "sicil"), "ad": ("ad soyad", "ad"), "departman": ("departman", "birim"), "pozisyon": ("pozisyon", "unvan"),
            "yonetici": ("yonetici", "bagli oldugu yonetici", "amir"), "giris": ("ise giris", "ise giris tarihi", "giris tarihi"),
            "cikis": ("ayrilis tarihi", "cikis tarihi", "isten cikis tarihi", "ayrilis"), "neden": ("ayrilis nedeni", "cikis nedeni", "neden"),
            "tur": ("ayrilis turu", "cikis turu")}
TURLER = ("Gönüllü", "Gönülsüz", "Diğer", "Belirsiz")
KIDEM_BANTLARI = [(0, 90, "0–3 ay"), (90, 365, "3–12 ay"), (365, 3 * 365, "1–3 yıl"), (3 * 365, 5 * 365, "3–5 yıl"), (5 * 365, 10 ** 6, "5 yıl +")]
# (anahtar kelime, tür) — sırayla denenir; ilk eşleşen kazanır
NEDEN_KURALLARI = [("isci tarafindan", "Gönüllü"), ("iscinin", "Gönüllü"), ("kendi istegi", "Gönüllü"),
                   ("emekli", "Diğer"), ("olum", "Diğer"), ("vefat", "Diğer"), ("askerlik", "Diğer"), ("malul", "Diğer"),
                   ("belirli sureli", "Diğer"), ("sozlesme sonu", "Diğer"), ("sozlesmenin sona", "Diğer"),
                   ("isveren", "Gönülsüz"), ("isten cikar", "Gönülsüz"), ("performans", "Gönülsüz"), ("disiplin", "Gönülsüz"),
                   ("toplu", "Gönülsüz"), ("kapan", "Gönülsüz"), ("fesih", "Gönülsüz"),
                   ("istifa", "Gönüllü"), ("evlilik", "Gönüllü")]


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def yuzde(x: float | None) -> str:
    return "—" if x is None else f"%{x * 100:.1f}".replace(".", ",")


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Kisi:
    sicil: str
    ad: str
    departman: str
    pozisyon: str
    yonetici: str
    giris: date
    cikis: date | None
    neden: str
    tur: str = ""

    def aktif(self, gun_sonu: date) -> bool:
        return self.giris <= gun_sonu and (self.cikis is None or self.cikis > gun_sonu)

    def kidem_gun(self, d: date) -> int:
        return (d - self.giris).days + 1


def tur_bul(neden: str, acik_tur: str = "") -> str:
    t = katla(acik_tur)
    if t:
        if t.startswith("gonullu"):
            return "Gönüllü"
        if t.startswith(("gonulsuz", "zorunlu")):
            return "Gönülsüz"
        if t.startswith("diger"):
            return "Diğer"
    n = katla(neden)
    return next((tur for k, tur in NEDEN_KURALLARI if k in n), "Belirsiz")


def oku_personel(yol: Path) -> tuple[list[Kisi], list[dict]]:
    kisiler, uyarilar = [], []
    for r in kayitlar(yol, SUTUNLAR, ("sicil", "giris")):
        if not r.get("sicil"):
            continue
        g, c = tarih(r.get("giris")), tarih(r.get("cikis"))
        if g is None:
            uyarilar.append({"onem": "Orta", "tur": "Okunamayan satır", "kim": metin(r["sicil"]), "aciklama": f"Satır {r['_satir']}: işe giriş tarihi yok"})
            continue
        if c and c < g:
            uyarilar.append({"onem": "Orta", "tur": "Tarih hatası", "kim": metin(r["sicil"]), "aciklama": f"Ayrılış ({c:%d.%m.%Y}) girişten önce; satır atlandı"})
            continue
        k = Kisi(metin(r["sicil"]), metin(r.get("ad")), metin(r.get("departman")) or "—", metin(r.get("pozisyon")), metin(r.get("yonetici")) or "—", g, c,
                 metin(r.get("neden")))
        if c:
            k.tur = tur_bul(k.neden, metin(r.get("tur")))
        kisiler.append(k)
    return kisiler, uyarilar


def kidem_bandi(gun: int) -> str:
    return next(ad for a, b, ad in KIDEM_BANTLARI if a <= gun < b)


def ay_sonlari(bas: date, bit: date) -> list[date]:
    sonuc, d = [], bas
    while True:
        sonraki = date(d.year + (d.month == 12), d.month % 12 + 1, 1)
        son = sonraki - timedelta(days=1)
        if son >= bit:
            sonuc.append(bit)
            return sonuc
        sonuc.append(son)
        d = sonraki


# ----------------------------------------------------------------------------
# Analiz
# ----------------------------------------------------------------------------

def analiz_et(kisiler: list[Kisi], bas: date, bit: date, erken_gun: int = 90, kat: float = 1.5, min_ekip: float = 3) -> dict:
    uyarilar = []
    donem_gun = (bit - bas).days + 1
    yillik = 365 / donem_gun
    olcum = [bas - timedelta(days=1)] + ay_sonlari(bas, bit)
    ayrilanlar = [k for k in kisiler if k.cikis and bas <= k.cikis <= bit]
    girenler = [k for k in kisiler if bas <= k.giris <= bit]

    def ortalama(filtre) -> float:
        return sum(sum(1 for k in kisiler if k.aktif(d) and filtre(k, d)) for d in olcum) / len(olcum)

    def oranlar(ort: float, lst: list[Kisi]) -> dict:
        say = Counter(k.tur for k in lst)
        x = {"ort": ort, "ayrilan": len(lst), **{t: say[t] for t in TURLER}}
        x["oran"] = len(lst) / ort if ort else None
        x["yillik"] = x["oran"] * yillik if ort else None
        x["gonullu_yillik"] = say["Gönüllü"] / ort * yillik if ort else None
        return x

    genel = oranlar(ortalama(lambda k, d: True), ayrilanlar)
    genel["bas"] = sum(1 for k in kisiler if k.aktif(olcum[0]))
    genel["son"] = sum(1 for k in kisiler if k.aktif(bit))
    genel["giren"] = len(girenler)
    kalan = [k for k in kisiler if k.aktif(olcum[0]) and k.aktif(bit)]
    genel["tutma"] = len(kalan) / genel["bas"] if genel["bas"] else None

    def kirilim(anahtar) -> list[dict]:
        sonuc = []
        for ad in sorted({anahtar(k) for k in kisiler}):
            x = oranlar(ortalama(lambda k, d, ad=ad: anahtar(k) == ad), [k for k in ayrilanlar if anahtar(k) == ad])
            if x["ort"] or x["ayrilan"]:
                sonuc.append({"ad": ad, **x})
        return sonuc

    departman = kirilim(lambda k: k.departman)
    yonetici = kirilim(lambda k: k.yonetici)
    kidem = []
    for _, _, ad in KIDEM_BANTLARI:
        ort = ortalama(lambda k, d, ad=ad: kidem_bandi(k.kidem_gun(d)) == ad)
        lst = [k for k in ayrilanlar if kidem_bandi(k.kidem_gun(k.cikis)) == ad]
        kidem.append({"ad": ad, **oranlar(ort, lst)})
    # Aylık trend
    aylik = []
    for i, son in enumerate(olcum[1:], 1):
        ilk = olcum[i - 1] + timedelta(days=1)
        b, s = sum(1 for k in kisiler if k.aktif(olcum[i - 1])), sum(1 for k in kisiler if k.aktif(son))
        ay = [k for k in ayrilanlar if ilk <= k.cikis <= son]
        aylik.append({"ay": ilk, "bas": b, "giren": sum(1 for k in girenler if ilk <= k.giris <= son), "ayrilan": len(ay),
                      "gonullu": sum(1 for k in ay if k.tur == "Gönüllü"), "son": s, "oran": len(ay) / ((b + s) / 2) if b + s else None})
    # Erken ayrılma
    ilk_yil = [k for k in ayrilanlar if k.kidem_gun(k.cikis) <= 365]
    kohort = [k for k in girenler if k.giris + timedelta(days=erken_gun) <= bit]
    erken = [k for k in kohort if k.cikis and (k.cikis - k.giris).days < erken_gun]
    erken_yon = []
    for y in sorted({k.yonetici for k in kohort}):
        ko = [k for k in kohort if k.yonetici == y]
        e = [k for k in erken if k.yonetici == y]
        erken_yon.append({"ad": y, "alinan": len(ko), "erken": len(e), "oran": len(e) / len(ko)})
        if len(e) >= 2:
            uyarilar.append({"onem": "Yüksek", "tur": "Erken ayrılma", "kim": y, "aciklama": f"Dönemde işe alınan {len(ko)} kişiden {len(e)}'si "
                             f"{erken_gun} gün içinde ayrıldı ({', '.join(k.sicil for k in e)}). İşe alım uyumu, oryantasyon ve ilk hafta deneyimi "
                             "gözden geçirilmeli; çıkış görüşmesi notlarına bakın."})
    # Sinyaller
    g_oran = genel["gonullu_yillik"] or 0
    for etiket, liste in (("Departman", departman), ("Yönetici", yonetici)):
        for x in liste:
            if x["ort"] >= min_ekip and x["Gönüllü"] >= 2 and x["gonullu_yillik"] > g_oran * kat:
                uyarilar.append({"onem": "Orta", "tur": f"Yüksek gönüllü devir ({etiket.lower()})", "kim": x["ad"],
                                 "aciklama": f"Yıllıklandırılmış gönüllü devir {yuzde(x['gonullu_yillik'])}, şirket {yuzde(g_oran)} "
                                 f"({x['Gönüllü']} gönüllü ayrılış, ortalama {x['ort']:.1f} kişi)"})
            elif x["ort"] < min_ekip and x["ayrilan"]:
                uyarilar.append({"onem": "Bilgi", "tur": "Küçük grup", "kim": x["ad"], "aciklama": f"{etiket} ortalaması {x['ort']:.1f} kişi; oranlar "
                                 "küçük sayılardan etkilenir, karşılaştırmada dikkatli olun"})
    for k in ayrilanlar:
        if k.tur == "Belirsiz":
            uyarilar.append({"onem": "Bilgi", "tur": "Ayrılış türü belirsiz", "kim": k.sicil, "aciklama": f"Neden '{k.neden or '—'}' sınıflanamadı; "
                             "'Ayrılış Türü' sütununa Gönüllü / Gönülsüz / Diğer yazın"})
    if ilk_yil and len(ilk_yil) / len(ayrilanlar) >= 0.3:
        uyarilar.append({"onem": "Orta", "tur": "İlk yıl ayrılışları", "kim": "", "aciklama": f"Ayrılanların {len(ilk_yil)}/{len(ayrilanlar)}'i "
                         f"({yuzde(len(ilk_yil) / len(ayrilanlar))}) ilk yılını doldurmadan ayrıldı"})
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    uyarilar.sort(key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))
    return {"genel": genel, "departman": departman, "yonetici": yonetici, "kidem": kidem, "aylik": aylik, "ayrilanlar": ayrilanlar,
            "ilk_yil": ilk_yil, "kohort": kohort, "erken": erken, "erken_yon": erken_yon, "uyarilar": uyarilar, "donem_gun": donem_gun,
            "olcum": olcum, "erken_gun": erken_gun}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
TUR_RENK = {"Gönüllü": "FFF4CE", "Gönülsüz": "FDE2E1", "Diğer": "E8F0FE", "Belirsiz": "EEEEEE"}
KONTROL = PatternFill("solid", fgColor="FFF4CE")
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, basliklar, genislik=()):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def _kirilim_yaz(ws, baslik, liste):
    _baslik(ws, [baslik, "Ort. Personel", "Ayrılan", "Gönüllü", "Gönülsüz", "Diğer", "Belirsiz", "Devir Oranı (dönem)", "Yıllıklandırılmış",
                 "Gönüllü (yıllık)"])
    for x in liste:
        ws.append([x["ad"], round(x["ort"], 1), x["ayrilan"], x["Gönüllü"], x["Gönülsüz"], x["Diğer"], x["Belirsiz"], x["oran"], x["yillik"], x["gonullu_yillik"]])
        for j in (8, 9, 10):
            ws.cell(ws.max_row, j).number_format = "0.0%"
    ws.append([])


def rapor_yaz(cikti: Path, s: dict, bas: date, bit: date) -> None:
    wb = Workbook()
    oz = wb.active
    oz.title = "Özet"
    g = s["genel"]
    _baslik(oz, ["Gösterge", "Değer", "Açıklama"], (34, 14, 80))
    satirlar = [("Dönem", f"{bas:%d.%m.%Y} – {bit:%d.%m.%Y}", f"{s['donem_gun']} gün"),
                ("Dönem başı personel", g["bas"], f"{s['olcum'][0]:%d.%m.%Y} gün sonu aktif"),
                ("Dönem sonu personel", g["son"], f"{bit:%d.%m.%Y} gün sonu aktif"),
                ("Ortalama personel", round(g["ort"], 2), f"Dönem başı ve {len(s['olcum']) - 1} ay sonu ölçümünün ortalaması"),
                ("İşe alınan", g["giren"], ""), ("Ayrılan", g["ayrilan"], ", ".join(f"{t} {g[t]}" for t in TURLER if g[t])),
                ("Devir oranı (dönem)", g["oran"], "Ayrılan ÷ ortalama personel"),
                ("Devir oranı (yıllıklandırılmış)", g["yillik"], "× 365 ÷ dönem günü"),
                ("Gönüllü devir (yıllıklandırılmış)", g["gonullu_yillik"], "Yalnız gönüllü ayrılışlar"),
                ("Tutma oranı", g["tutma"], "Dönem başındaki personelden dönem sonunda hâlâ çalışanların oranı"),
                ("İlk yılında ayrılan", len(s["ilk_yil"]), f"Ayrılanların {yuzde(len(s['ilk_yil']) / g['ayrilan']) if g['ayrilan'] else '—'}'i"),
                (f"{s['erken_gun']} gün içinde ayrılan (yeni alım)", f"{len(s['erken'])} / {len(s['kohort'])}",
                 f"Dönemde işe alınıp {s['erken_gun']} günü gözlenebilen {len(s['kohort'])} kişiden")]
    for ad, v, a in satirlar:
        oz.append([ad, v, a])
        if isinstance(v, float) and ad.startswith(("Devir", "Gönüllü", "Tutma")):
            oz.cell(oz.max_row, 2).number_format = "0.0%"
    oz.append([])
    oz.append(["Not", "", "Ayrılış türü 'Ayrılış Türü' sütunundan, yoksa nedenden çıkarılır. Oranlar küçük gruplarda (ortalama < 3 kişi) yanıltıcı olabilir."])

    kr = wb.create_sheet("Kırılımlar")
    _kirilim_yaz(kr, "Departman", s["departman"])
    _kirilim_yaz(kr, "Yönetici", s["yonetici"])
    _kirilim_yaz(kr, "Kıdem (ayrılışta)", s["kidem"])
    for j, w in enumerate((22, 12, 9, 9, 9, 8, 9, 12, 12, 12), 1):
        kr.column_dimensions[get_column_letter(j)].width = w

    ay = wb.create_sheet("Aylık Trend")
    _baslik(ay, ["Ay", "Ay Başı", "Giren", "Ayrılan", "Gönüllü", "Ay Sonu", "Aylık Devir"], (10, 9, 8, 9, 9, 9, 11))
    for x in s["aylik"]:
        ay.append([f"{x['ay']:%m.%Y}", x["bas"], x["giren"], x["ayrilan"], x["gonullu"], x["son"], x["oran"]])
        ay.cell(ay.max_row, 7).number_format = "0.0%"
    if s["aylik"]:
        gr = LineChart()
        gr.title, gr.height, gr.width = "Aylık giren / ayrılan", 7, 16
        gr.add_data(Reference(ay, min_col=3, max_col=4, min_row=1, max_row=1 + len(s["aylik"])), titles_from_data=True)
        gr.set_categories(Reference(ay, min_col=1, min_row=2, max_row=1 + len(s["aylik"])))
        ay.add_chart(gr, "I2")

    er = wb.create_sheet("Erken Ayrılma")
    _baslik(er, ["Yönetici", "Dönemde İşe Alınan (gözlenebilir)", f"{s['erken_gun']} Gün İçinde Ayrılan", "Oran"], (18, 18, 16, 9))
    for x in s["erken_yon"]:
        er.append([x["ad"], x["alinan"], x["erken"], x["oran"]])
        er.cell(er.max_row, 4).number_format = "0.0%"
        if x["erken"] >= 2:
            er.cell(er.max_row, 3).fill = PatternFill("solid", fgColor="FDE2E1")
    er.append([])
    _baslik(er, ["Sicil", "Ad Soyad", "Departman", "Yönetici", "İşe Giriş", "Ayrılış", "Gün", "Neden"])
    for k in s["erken"]:
        er.append([k.sicil, k.ad, k.departman, k.yonetici, k.giris, k.cikis, (k.cikis - k.giris).days + 1, k.neden])
        er.cell(er.max_row, 5).number_format = er.cell(er.max_row, 6).number_format = "DD.MM.YYYY"

    al = wb.create_sheet("Ayrılanlar")
    _baslik(al, ["Sicil", "Ad Soyad", "Departman", "Pozisyon", "Yönetici", "İşe Giriş", "Ayrılış", "Kıdem (gün)", "Kıdem Bandı", "Neden", "Tür",
                 "Çıkış Görüşmesi Notu"], (8, 16, 12, 18, 14, 11, 11, 9, 10, 34, 10, 30))
    for k in sorted(s["ayrilanlar"], key=lambda k: k.cikis):
        gun = k.kidem_gun(k.cikis)
        al.append([k.sicil, k.ad, k.departman, k.pozisyon, k.yonetici, k.giris, k.cikis, gun, kidem_bandi(gun), k.neden, k.tur, ""])
        r = al.max_row
        al.cell(r, 6).number_format = al.cell(r, 7).number_format = "DD.MM.YYYY"
        al.cell(r, 11).fill = PatternFill("solid", fgColor=TUR_RENK[k.tur])
        al.cell(r, 12).fill = KONTROL
    al.auto_filter.ref = f"A1:L{al.max_row}"
    al.freeze_panes = "A2"

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Kim", "Açıklama", "Aksiyon"], (9, 28, 14, 90, 24))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["kim"], u["aciklama"], ""])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).alignment = UST
        uy.cell(uy.max_row, 5).fill = KONTROL
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(yol: Path, cikti: Path, bas: date, bit: date, erken_gun: int = 90, kat: float = 1.5) -> dict:
    kisiler, okuma = oku_personel(yol)
    if not kisiler:
        raise ValueError(f"{yol.name}: personel kaydı bulunamadı")
    s = analiz_et(kisiler, bas, bit, erken_gun, kat)
    s["uyarilar"] = okuma + s["uyarilar"]
    rapor_yaz(cikti, s, bas, bit)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Giriş-çıkış verisinden personel devir oranlarını ve erken ayrılma sinyallerini çıkarır.")
    p.add_argument("--personel", type=Path, default=ORNEK / "personel.csv",
                   help="Sicil, Ad, Departman, Pozisyon, Yönetici, İşe Giriş, Ayrılış Tarihi, Ayrılış Nedeni, [Ayrılış Türü]")
    p.add_argument("--donem", nargs=2, metavar=("BAŞLANGIÇ", "BİTİŞ"), help="Analiz dönemi GG.AA.YYYY GG.AA.YYYY")
    p.add_argument("--erken-gun", type=int, default=90, help="Erken ayrılma eşiği, gün (varsayılan 90)")
    p.add_argument("--kat", type=float, default=1.5, help="Şirket gönüllü devrinin kaç katı 'yüksek' sayılır (varsayılan 1,5)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "personel_devir_orani.xlsx")
    a = p.parse_args(argv)
    if not a.personel.exists():
        print(f"[X] Dosya bulunamadı: {a.personel}")
        return 1
    if a.donem:
        bas, bit = tarih(a.donem[0]), tarih(a.donem[1])
        if not bas or not bit or bas > bit:
            print("[X] --donem iki geçerli tarih olmalı: GG.AA.YYYY GG.AA.YYYY")
            return 2
    elif a.personel == ORNEK / "personel.csv":
        bas, bit = date(2026, 1, 1), date(2026, 9, 30)
    else:
        bugun = date.today()
        bas, bit = date(bugun.year - 1, bugun.month, 1), date(bugun.year, bugun.month, 1) - timedelta(days=1)
        print(f"[i] --donem verilmedi; son 12 tam ay kullanıldı: {bas:%d.%m.%Y} – {bit:%d.%m.%Y}")
    try:
        s = calistir(a.personel, a.cikti, bas, bit, a.erken_gun, a.kat)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    g = s["genel"]
    print(f"[OK] Ortalama personel {g['ort']:.1f} · ayrılan {g['ayrilan']} (gönüllü {g['Gönüllü']}) · devir {yuzde(g['oran'])} "
          f"(yıllık {yuzde(g['yillik'])}, gönüllü yıllık {yuzde(g['gonullu_yillik'])})")
    for u in s["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['kim']} {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
