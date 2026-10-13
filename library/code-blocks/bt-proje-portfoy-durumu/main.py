"""
BT Proje Portföy Durumu — Workers / Workless kod bloğu
Bilgi Teknolojileri › BT Müdürü

Projelerin takvim, bütçe ve risk durumunu kazanılmış değer yöntemiyle tek yönetim tablosunda toplar.
  Planlanan ilerleme % = (--tarih − başlangıç) ÷ (planlanan bitiş − başlangıç), 0–100 arasında sınırlanır
  Planlanan değer (PD) = bütçe × planlanan ilerleme %;   Kazanılan değer (KD) = bütçe × gerçekleşen ilerleme %
  Takvim endeksi (TPE) = KD ÷ PD;   Maliyet endeksi (MPE) = KD ÷ harcanan      (1'in altı: geride / pahalı)
  Tahmini toplam maliyet = bütçe ÷ MPE;   tahmini aşım = tahmini toplam maliyet − bütçe
  Renk                 = endeks ≥ --yesil (0,95) Yeşil · ≥ --sari (0,85) Sarı · altı Kırmızı
                         Takvim: planlanan bitişi geçmiş ve bitmemiş proje Kırmızı; tahmini bitişi planı aşan proje en az Sarı
                         Bütçe: harcaması bütçeyi aşmış proje Kırmızı
  Genel                = takvim ve bütçe renginin kötü olanı. Risk sütunu sizin değerlendirmenizdir, ayrıca gösterilir
Yalnız durumu "Devam" olan projeler renklendirilir. İlerleme yüzdesini siz verirsiniz; kod doğruluğunu sınayamaz.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
SUTUN = {"proje": ("proje", "proje adi"), "sahip": ("sahip", "proje yoneticisi", "sorumlu", "sponsor"), "bas": ("baslangic", "baslangic tarihi"), "bitis": ("planlanan bitis", "bitis", "hedef bitis", "planlanan bitis tarihi"),
         "tahmini": ("tahmini bitis", "guncel bitis", "beklenen bitis"), "butce": ("butce", "onayli butce"), "harcanan": ("harcanan", "gerceklesen", "fiili harcama", "harcama"),
         "ilerleme": ("ilerleme", "ilerleme yuzdesi", "tamamlanma", "tamamlanma yuzdesi"), "risk": ("risk", "risk seviyesi"), "durum": ("durum",), "oncelik": ("oncelik",)}
DURUM = {"": "Devam", "devam": "Devam", "devam ediyor": "Devam", "aktif": "Devam", "tamamlandi": "Tamamlandı", "bitti": "Tamamlandı", "kapandi": "Tamamlandı", "beklemede": "Beklemede", "durduruldu": "Beklemede", "askida": "Beklemede",
         "planlandi": "Beklemede", "baslamadi": "Beklemede", "iptal": "İptal", "iptal edildi": "İptal"}
RISK = {"yuksek": "Yüksek", "kritik": "Yüksek", "orta": "Orta", "dusuk": "Düşük", "": ""}
RENK_AD = ("Yeşil", "Sarı", "Kırmızı")
R_DOLGU = {"Yeşil": YESIL, "Sarı": SARI, "Kırmızı": KIRMIZI, "Yüksek": KIRMIZI, "Orta": SARI, "Düşük": YESIL}
Z, YUZ = Decimal(0), Decimal(100)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("proje", "bas", "bitis", "ilerleme")):
        ad = metin(r["proje"])
        if not ad:
            continue
        bas, bit, il, d = tarih(r["bas"]), tarih(r["bitis"]), para(r["ilerleme"]), katla(r["durum"])
        if bas is None or bit is None or bit <= bas or il is None or not 0 <= il <= 100:
            uy.append(uyari("Yüksek", "Proje okunamadı", f"{ad} (satır {r['_satir']})", f"Başlangıç '{metin(r['bas'])}', planlanan bitiş '{metin(r['bitis'])}', ilerleme '{metin(r['ilerleme'])}'. Bitiş başlangıçtan sonra, ilerleme 0–100 arasında olmalı. "
                                                                                    "Proje rapora alınmadı."))
            continue
        if d not in DURUM:
            uy.append(uyari("Orta", "Durum tanınmadı", ad, f"'{metin(r['durum'])}' tanınmadı; proje 'Devam' sayıldı. Devam, Tamamlandı, Beklemede veya İptal yazın."))
        butce, harcanan = para(r["butce"]), para(r["harcanan"])
        l.append({"proje": ad, "sahip": metin(r["sahip"]), "bas": bas, "bitis": bit, "tahmini": tarih(r["tahmini"]), "butce": butce if butce and butce > 0 else None, "harcanan": harcanan if harcanan and harcanan > 0 else Z, "ilerleme": il,
                  "risk": RISK.get(katla(r["risk"]), metin(r["risk"])), "durum": DURUM.get(d, "Devam"), "oncelik": metin(r["oncelik"])})
    return l, uy


def renk(endeks: Decimal | None, yesil: Decimal, sari: Decimal) -> str | None:
    return None if endeks is None else "Yeşil" if endeks >= yesil else "Sarı" if endeks >= sari else "Kırmızı"


def kotu(*renkler) -> str | None:
    r = [x for x in renkler if x]
    return max(r, key=RENK_AD.index) if r else None


def degerlendir(projeler: list[dict], bugun: date, yesil: Decimal = Decimal("0.95"), sari: Decimal = Decimal("0.85")) -> dict:
    if not projeler:
        raise ValueError("Raporlanacak proje yok.")
    uy = []
    for p in projeler:
        plan_oran = min(max(Decimal((bugun - p["bas"]).days) / (p["bitis"] - p["bas"]).days, Z), Decimal(1))
        p["plan_yuzde"] = yuvarla(plan_oran * YUZ)
        p["tpe_ham"] = p["ilerleme"] / YUZ / plan_oran if plan_oran else None
        p["pd"], p["kd"] = (yuvarla(p["butce"] * plan_oran), yuvarla(p["butce"] * p["ilerleme"] / YUZ)) if p["butce"] else (None, None)
        p["mpe_ham"] = p["butce"] * p["ilerleme"] / YUZ / p["harcanan"] if p["butce"] and p["harcanan"] and p["ilerleme"] else None      # ilerleme sıfırken endeks anlamsızdır
        p["tpe"], p["mpe"] = (yuvarla(x) if x is not None else None for x in (p["tpe_ham"], p["mpe_ham"]))
        p["tahmini_maliyet"] = yuvarla(p["butce"] / p["mpe_ham"]) if p["mpe_ham"] else None
        p["asim"] = p["tahmini_maliyet"] - p["butce"] if p["tahmini_maliyet"] is not None else None
        p["gecikme"] = max((p["tahmini"] - p["bitis"]).days, 0) if p["tahmini"] else 0
        p["takvim"] = p["butce_renk"] = p["genel"] = None
        if p["durum"] == "Tamamlandı" and p["ilerleme"] < 100:
            uy.append(uyari("Orta", "Tamamlandı ama ilerleme eksik", p["proje"], f"Durum 'Tamamlandı', ilerleme %{p['ilerleme']}. Kaydı düzeltin."))
        if p["durum"] != "Devam":
            continue
        bitmis = p["ilerleme"] >= 100
        p["takvim"] = "Yeşil" if bitmis else renk(p["tpe_ham"], yesil, sari) or "Yeşil"
        if not bitmis and p["bitis"] < bugun:
            p["takvim"] = "Kırmızı"
            uy.append(uyari("Yüksek", "Planlanan bitiş geçmiş", p["proje"], f"Planlanan bitiş {p['bitis']:%d.%m.%Y}, {(bugun - p['bitis']).days} gün geçti; ilerleme %{p['ilerleme']}."
                            + (f" Güncel tahmin {p['tahmini']:%d.%m.%Y}." if p["tahmini"] else " Güncel bitiş tahmini yazılmamış.")))
        elif p["gecikme"]:
            p["takvim"] = kotu(p["takvim"], "Sarı")
            uy.append(uyari("Orta", "Tahmini bitiş planı aşıyor", p["proje"], f"Planlanan bitiş {p['bitis']:%d.%m.%Y}, tahmini bitiş {p['tahmini']:%d.%m.%Y}; {p['gecikme']} gün gecikme bekleniyor."))
        if bitmis:
            uy.append(uyari("Bilgi", "İlerleme %100 ama proje açık", p["proje"], "İlerleme tamamlanmış görünüyor; kapanış işlemleri bittiyse durumu 'Tamamlandı' yapın."))
        if p["butce"] is None:
            uy.append(uyari("Bilgi", "Bütçe yok", p["proje"], "Bütçe yazılmamış; maliyet endeksi ve bütçe rengi hesaplanmadı."))
        else:
            p["butce_renk"] = renk(p["mpe_ham"], yesil, sari) or "Yeşil"
            if p["harcanan"] > p["butce"]:
                p["butce_renk"] = "Kırmızı"
                uy.append(uyari("Yüksek", "Bütçe aşıldı", p["proje"], f"Harcanan {p['harcanan']:,.2f}, bütçe {p['butce']:,.2f}; aşım {p['harcanan'] - p['butce']:,.2f}."))
            elif p["harcanan"] and p["ilerleme"] == 0:
                p["butce_renk"] = kotu(p["butce_renk"], "Sarı")
                uy.append(uyari("Orta", "Harcama var, ilerleme yok", p["proje"], f"{p['harcanan']:,.2f} harcanmış ama ilerleme %0. İlerleme kaydını güncelleyin."))
        p["genel"] = kotu(p["takvim"], p["butce_renk"])
        if p["genel"] == "Kırmızı":
            uy.append(uyari("Yüksek", "Kırmızı proje", p["proje"], f"Takvim {p['takvim']}" + (f" (TPE {p['tpe']})" if p["tpe"] is not None else "") + f", bütçe {p['butce_renk'] or '-'}" + (f" (MPE {p['mpe']})" if p["mpe"] is not None else "")
                            + (f"; bu verimle tahmini toplam maliyet {p['tahmini_maliyet']:,.2f} (aşım {p['asim']:,.2f})." if p["asim"] and p["asim"] > 0 else ".")))
    aktif = [p for p in projeler if p["durum"] == "Devam"]
    top = {k: sum((p[k] for p in aktif if p["butce"] and p[k] is not None), Z) for k in ("butce", "harcanan", "kd", "tahmini_maliyet")}
    top["tahmini_maliyet"] += sum((p["butce"] for p in aktif if p["butce"] and p["tahmini_maliyet"] is None), Z)     # harcaması olmayan proje bütçesiyle sayılır
    sira = {"Kırmızı": 0, "Sarı": 1, "Yeşil": 2, None: 3}
    projeler.sort(key=lambda p: (p["durum"] != "Devam", sira[p["genel"]], katla(p["oncelik"]) or "~", katla(p["proje"])))
    return {"projeler": projeler, "aktif": len(aktif), "renkler": {r: sum(p["genel"] == r for p in aktif) for r in reversed(RENK_AD)}, "durumlar": {d: sum(p["durum"] == d for p in projeler) for d in ("Devam", "Tamamlandı", "Beklemede", "İptal")},
            "toplam": top, "mpe": yuvarla(top["kd"] / top["harcanan"]) if top["harcanan"] else None, "bugun": bugun, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Portföy"
    _baslik(ws, ["Genel", "Proje", "Sahip", "Öncelik", "Durum", "Başlangıç", "Planlanan Bitiş", "Tahmini Bitiş", "Gecikme (gün)", "Planlanan %", "Gerçekleşen %", "TPE", "Takvim", "Bütçe", "Harcanan", "Kazanılan Değer", "MPE", "Bütçe Durumu",
                 "Tahmini Toplam Maliyet", "Tahmini Aşım", "Risk"], (9, 34, 18, 8, 12, 12, 12, 12, 9, 10, 10, 7, 9, 16, 16, 16, 7, 9, 17, 16, 9))
    for p in r["projeler"]:
        _satir(ws, [p["genel"] or "-", p["proje"], p["sahip"], p["oncelik"], p["durum"], p["bas"], p["bitis"], p["tahmini"], p["gecikme"] or None, p["plan_yuzde"], p["ilerleme"], p["tpe"], p["takvim"], p["butce"], p["harcanan"] or None,
                    p["kd"], p["mpe"], p["butce_renk"], p["tahmini_maliyet"], p["asim"], p["risk"]], para_s=(10, 11, 12, 14, 15, 16, 17, 19, 20), tarih_s=(6, 7, 8),
               dolgu={1: R_DOLGU.get(p["genel"]), 13: R_DOLGU.get(p["takvim"]), 18: R_DOLGU.get(p["butce_renk"]), 21: R_DOLGU.get(p["risk"])})
    ws = wb.create_sheet("Özet")
    _baslik(ws, ["Gösterge", "Değer"], (44, 18))
    t = r["toplam"]
    for etiket, deger in [(f"{d} proje", n) for d, n in r["durumlar"].items() if n] + [(f"Devam edenlerde {k}", n) for k, n in r["renkler"].items()]:
        _satir(ws, [etiket, deger])
    for etiket, deger in (("Devam eden projelerin bütçesi", t["butce"]), ("Harcanan", t["harcanan"]), ("Kazanılan değer", t["kd"]), ("Portföy maliyet endeksi (MPE)", r["mpe"]), ("Tahmini toplam maliyet", t["tahmini_maliyet"]),
                          ("Tahmini aşım (+) / tasarruf (−)", t["tahmini_maliyet"] - t["butce"])):
        _satir(ws, [etiket, deger], para_s=(2,))
    ws.append([])
    ws.append([f"Durum tarihi {r['bugun']:%d.%m.%Y}. TPE = kazanılan değer ÷ planlanan değer; MPE = kazanılan değer ÷ harcanan. Tahmini toplam maliyet, bugünkü verimin süreceği varsayımına dayanır."])
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, yesil: Decimal = Decimal("0.95"), sari: Decimal = Decimal("0.85")) -> dict:
    projeler, uy = oku(yol)
    r = degerlendir(projeler, bugun, yesil, sari)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="BT projelerinin takvim, bütçe ve risk durumunu tek tabloda toplar.")
    p.add_argument("--projeler", type=Path, help="Proje, Sahip, Başlangıç, Planlanan Bitiş, Tahmini Bitiş, Bütçe, Harcanan, İlerleme %%, Risk, Durum, Öncelik")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--yesil", default="0,95", help="Endeks bu değer ve üstündeyse Yeşil (varsayılan 0,95)")
    p.add_argument("--sari", default="0,85", help="Endeks bu değer ve üstündeyse Sarı, altındaysa Kırmızı (varsayılan 0,85)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "bt_proje_portfoy.xlsx")
    a = p.parse_args(argv)
    ornek = a.projeler is None
    yol = ORNEK / "projeler.csv" if ornek else a.projeler
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    yesil, sari = para(a.yesil), para(a.sari)
    if bugun is None or yesil is None or sari is None or not 0 < sari < yesil <= 1:
        print("[X] --tarih GG.AA.YYYY biçiminde olmalı; --sari sıfırdan büyük ve --yesil değerinden küçük, --yesil en çok 1 olmalı (örn. 0,95 ve 0,85)")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, bugun, yesil, sari)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    k = r["renkler"]
    print(f"[OK] {len(r['projeler'])} proje, {r['aktif']} devam ediyor · Kırmızı {k['Kırmızı']}, Sarı {k['Sarı']}, Yeşil {k['Yeşil']} · bütçe {r['toplam']['butce']:,.2f}, harcanan {r['toplam']['harcanan']:,.2f}"
          + (f", portföy MPE {r['mpe']}" if r["mpe"] is not None else ""))
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
