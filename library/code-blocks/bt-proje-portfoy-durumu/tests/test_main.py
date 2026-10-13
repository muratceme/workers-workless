import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)


def p(ad, ilerleme, butce=1000, harcanan=0, bas=date(2026, 9, 2), bitis=date(2026, 12, 11), tahmini=None, durum="Devam", oncelik="", risk=""):
    """Varsayılan takvim: 100 günlük proje, durum tarihinde 40. gün → planlanan ilerleme %40."""
    return {"proje": ad, "sahip": "", "bas": bas, "bitis": bitis, "tahmini": tahmini, "butce": None if butce is None else D(butce), "harcanan": D(harcanan), "ilerleme": D(str(ilerleme)), "risk": risk, "durum": durum, "oncelik": oncelik}


def tek(*a, **kw):
    r = main.degerlendir([p("P", *a, **kw)], BUGUN)
    return r["projeler"][0], sorted(u["tur"] for u in r["uyarilar"])


class HesapTesti(unittest.TestCase):
    def test_kazanilmis_deger(self):
        x, uy = tek(30, harcanan=400)
        self.assertEqual((x["plan_yuzde"], x["pd"], x["kd"], x["tpe"], x["mpe"], x["tahmini_maliyet"], x["asim"], x["takvim"], x["butce_renk"], x["genel"], uy),
                         (D("40.00"), D("400.00"), D("300.00"), D("0.75"), D("0.75"), D("1333.33"), D("333.33"), "Kırmızı", "Kırmızı", "Kırmızı", ["Kırmızı proje"]))
        x, uy = tek(50, harcanan=400)
        self.assertEqual((x["tpe"], x["mpe"], x["tahmini_maliyet"], x["asim"], x["genel"], uy), (D("1.25"), D("1.25"), D("800.00"), D("-200.00"), "Yeşil", []))
        self.assertEqual([tek(i)[0]["takvim"] for i in (38, 37.9, 34, 33.9)], ["Yeşil", "Sarı", "Sarı", "Kırmızı"])                       # TPE 0,95 · 0,9475 · 0,85 · 0,8475
        self.assertEqual([tek(40, harcanan=h)[0]["butce_renk"] for h in (0, 421, 422, 470, 471)], ["Yeşil", "Yeşil", "Sarı", "Sarı", "Kırmızı"])   # MPE 0,9501 · 0,9479 · 0,8511 · 0,8493
        self.assertEqual([(tek(10, bas=b, bitis=s)[0]["plan_yuzde"], tek(10, bas=b, bitis=s)[0]["tpe"]) for b, s in ((date(2026, 11, 1), date(2026, 12, 1)), (date(2026, 10, 12), date(2026, 10, 22)), (date(2026, 10, 2), date(2026, 10, 12)))],
                         [(D("0.00"), None), (D("0.00"), None), (D("100.00"), D("0.10"))])

    def test_ek_kurallar(self):
        x, uy = tek(90, bas=date(2026, 6, 1), bitis=date(2026, 9, 30), tahmini=date(2026, 10, 31))
        self.assertEqual((x["takvim"], x["tpe"], x["gecikme"], x["genel"], uy), ("Kırmızı", D("0.90"), 31, "Kırmızı", ["Kırmızı proje", "Planlanan bitiş geçmiş"]))
        x, uy = tek(40, tahmini=date(2026, 12, 21))
        self.assertEqual((x["takvim"], x["gecikme"], x["genel"], uy), ("Sarı", 10, "Sarı", ["Tahmini bitiş planı aşıyor"]))
        self.assertEqual(tek(40, tahmini=date(2026, 12, 1))[0]["gecikme"], 0)
        x, uy = tek(99, harcanan=1001)
        self.assertEqual((x["butce_renk"], x["mpe"], uy), ("Kırmızı", D("0.99"), ["Bütçe aşıldı", "Kırmızı proje"]))
        x, uy = tek(0, harcanan=50, bas=date(2026, 10, 12), bitis=date(2026, 12, 31))
        self.assertEqual((x["takvim"], x["butce_renk"], x["mpe"], x["tahmini_maliyet"], uy), ("Yeşil", "Sarı", None, None, ["Harcama var, ilerleme yok"]))
        x, uy = tek(40, butce=None, harcanan=50)
        self.assertEqual((x["pd"], x["kd"], x["mpe"], x["butce_renk"], x["genel"], uy), (None, None, None, None, "Yeşil", ["Bütçe yok"]))
        x, uy = tek(100, harcanan=900, bitis=date(2026, 10, 1))
        self.assertEqual((x["takvim"], x["genel"], uy), ("Yeşil", "Yeşil", ["İlerleme %100 ama proje açık"]))
        for durum, beklenen in (("Tamamlandı", ["Tamamlandı ama ilerleme eksik"]), ("Beklemede", []), ("İptal", [])):
            x, uy = tek(10, harcanan=2000, bitis=date(2026, 10, 1), durum=durum)
            self.assertEqual((x["genel"], x["takvim"], uy), (None, None, beklenen))

    def test_portfoy(self):
        l = [p("Yeşil-2", 50, 1000, 400, oncelik="2"), p("Kırmızı", 30, 1000, 400, oncelik="3"), p("Sarı", 36, 2000, 720), p("Yeşil-1", 40, 500, 0, oncelik="1"), p("Bitti", 100, 9000, 9000, durum="Tamamlandı"),
             p("Bütçesiz", 40, None, 70)]
        r = main.degerlendir(l, BUGUN)
        self.assertEqual([(x["proje"], x["genel"]) for x in r["projeler"]], [("Kırmızı", "Kırmızı"), ("Sarı", "Sarı"), ("Yeşil-1", "Yeşil"), ("Yeşil-2", "Yeşil"), ("Bütçesiz", "Yeşil"), ("Bitti", None)])
        self.assertEqual((r["aktif"], r["renkler"], r["durumlar"], r["toplam"], r["mpe"]),
                         (5, {"Kırmızı": 1, "Sarı": 1, "Yeşil": 3}, {"Devam": 5, "Tamamlandı": 1, "Beklemede": 0, "İptal": 0},
                          {"butce": D(4500), "harcanan": D(1520), "kd": D("1720.00"), "tahmini_maliyet": D("4633.33")}, D("1.13")))
        with self.assertRaises(ValueError):
            main.degerlendir([], BUGUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "p.xlsx"
            r = main.calistir(ORNEK / "projeler.csv", cikti, BUGUN)
            self.assertEqual([(x["proje"], x["genel"], x["plan_yuzde"], x["tpe"], x["takvim"], x["mpe"], x["butce_renk"], x["tahmini_maliyet"]) for x in r["projeler"]],
                             [("ERP Sürüm Yükseltme", "Kırmızı", D("70.80"), D("0.78"), "Kırmızı", D("0.71"), "Kırmızı", D("5636363.64")),
                              ("E-posta Sistemi Göçü", "Kırmızı", D("100.00"), D("0.90"), "Kırmızı", D("0.93"), "Sarı", D("644444.44")),
                              ("Ağ Anahtarları Yenileme", "Sarı", D("22.78"), D("0.88"), "Sarı", D("1.14"), "Yeşil", D("1750000.00")),
                              ("Yedekleme Altyapısı Yenileme", "Yeşil", D("67.76"), D("1.03"), "Yeşil", D("1.31"), "Yeşil", D("914285.71")),
                              ("Siber Güvenlik Farkındalık Eğitimi", None, D("100.00"), D("1.00"), None, D("1.04"), None, D("240000.00")),
                              ("Veri Ambarı Kurulumu", None, D("0.00"), None, None, None, None, None)])
            self.assertEqual((r["renkler"], r["toplam"], r["mpe"]),
                             ({"Kırmızı": 2, "Sarı": 1, "Yeşil": 1}, {"butce": D(7800000), "harcanan": D(4670000), "kd": D(3980000), "tahmini_maliyet": D("8945093.79")}, D("0.85")))
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Yüksek", "Kırmızı proje", "E-posta Sistemi Göçü"), ("Yüksek", "Kırmızı proje", "ERP Sürüm Yükseltme"), ("Yüksek", "Planlanan bitiş geçmiş", "E-posta Sistemi Göçü"),
                              ("Orta", "Tahmini bitiş planı aşıyor", "ERP Sürüm Yükseltme")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Portföy", "Özet", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "p.csv"
            y.write_text("Proje Adı;Proje Yöneticisi;Başlangıç Tarihi;Bitiş;Onaylı Bütçe;Gerçekleşen;Tamamlanma;Risk;Durum\nA;Ali;01.01.2026;31.12.2026;1.000.000;250.000,50;%25;KRİTİK;\nB;;01.01.2026;31.12.2026;;;10;;askıda\n"
                         "C;;01.01.2026;31.12.2026;;;10;;yolunda\nD;;01.01.2026;01.01.2026;;;10;;\nE;;01.01.2026;31.12.2026;;;101;;\nF;;x;31.12.2026;;;5;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["proje"], x["sahip"], x["butce"], x["harcanan"], x["ilerleme"], x["risk"], x["durum"]) for x in l], [u["tur"] for u in uy]),
                             ([("A", "Ali", D(1000000), D("250000.50"), D(25), "Yüksek", "Devam"), ("B", "", None, D(0), D(10), "", "Beklemede"), ("C", "", None, D(0), D(10), "", "Devam")],
                              ["Durum tanınmadı"] + ["Proje okunamadı"] * 3))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--projeler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "12.10.2026", "--yesil", "0,90", "--sari", "0.80"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + x) for x in (["--yesil", "0,80"], ["--yesil", "1,1"], ["--sari", "0"], ["--sari", "x"], ["--tarih", "x"])], [2] * 5)
            self.assertEqual(main.main(["--projeler", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Proje;Başlangıç;Bitiş\nA;01.01.2026;31.12.2026\n", "Proje;Başlangıç;Bitiş;İlerleme\nA;x;31.12.2026;5\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
