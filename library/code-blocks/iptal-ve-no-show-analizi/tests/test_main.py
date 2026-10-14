import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def r(no, rez, giris, gece, kanal, oda, durum, iptal=None, tutar=None):
    g = date(2026, *giris)
    rz = date(2026, *rez) if rez else None
    return {"no": no, "rez": rz, "giris": g, "gece": gece, "kanal": kanal, "oda": oda, "durum": durum, "iptal": date(2026, *iptal) if iptal else None,
            "tutar": None if tutar is None else D(tutar), "lead": (g - rz).days if rz else None}


def veri():
    return [r("R1", (9, 1), (10, 1), 2, "OTA", "STD", "Gerçekleşti", tutar=2000), r("R2", (9, 20), (10, 1), 1, "OTA", "STD", "İptal", (9, 30), 1000),
            r("R3", (9, 25), (10, 2), 3, "OTA", "DLX", "No-show", tutar=3000), r("R4", (7, 1), (10, 5), 3, "Direkt", "DLX", "İptal", (9, 1), 4500),
            r("R5", None, (10, 6), 1, "Direkt", "STD", "Gerçekleşti")]


def k(s):
    return (s["ad"], s["rez"], s["iptal"], s["noshow"], s["iptal_oran"], s["noshow_oran"])


class AnalizTesti(unittest.TestCase):
    def test_oranlar(self):
        x = main.analiz(veri())
        t = x["toplam"]
        self.assertEqual((t["rez"], t["iptal_oran"], t["noshow_oran"], t["gerceklesen"], t["gece"], t["kayip_gece"], t["kayip_gelir"], t["lead"]),
                         (5, D("40.00"), D("33.33"), 2, 10, 7, D(8500), D("36.00")))
        self.assertEqual([k(s) for s in x["kanal"]], [("OTA", 3, 1, 1, D("33.33"), D("50.00")), ("Direkt", 2, 1, 0, D("50.00"), D("0.00"))])
        self.assertEqual([(s["ad"], s["rez"]) for s in x["lead"]], [("0–7 gün", 1), ("8–30 gün", 2), ("90+ gün", 1)])
        self.assertEqual([(s["ad"], s["adet"], s["pay"]) for s in x["iptal_zaman"]],
                         [("Aynı gün", 0, D("0.00")), ("1–3 gün", 1, D("50.00")), ("4–7 gün", 0, D("0.00")), ("8–30 gün", 0, D("0.00")), ("30+ gün", 1, D("50.00"))])
        self.assertEqual([(u["tur"], u["kim"]) for u in x["uyarilar"]], [("Son dakika iptalleri", "1 iptal")])

    def test_esikler(self):
        x = main.analiz(veri(), 2, D(30), D(5))
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in x["uyarilar"] if u["onem"] == "Orta"),
                         [("No-show oranı yüksek", "Kanal: OTA"), ("No-show oranı yüksek", "Oda tipi: DLX"), ("İptal oranı yüksek", "Kanal: Direkt"), ("İptal oranı yüksek", "Kanal: OTA"),
                          ("İptal oranı yüksek", "Oda tipi: DLX"), ("İptal oranı yüksek", "Oda tipi: STD")])
        self.assertEqual([main._durum(s) for s in ("Konakladı", "İPTAL", "No Show", "noshow", "Gelmedi", "", "x")], ["Gerçekleşti", "İptal", "No-show", "No-show", "No-show", None, None])
        with self.assertRaises(ValueError):
            main.analiz([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            x = main.calistir(ORNEK / "rezervasyonlar.csv", cikti)
            self.assertEqual(([k(s) for s in x["kanal"]], k({"ad": "T", **x["toplam"]})),
                             ([("OTA A", 84, 29, 5, D("34.52"), D("9.09")), ("Web sitesi", 52, 2, 0, D("3.85"), D("0.00")), ("Tur operatörü", 46, 4, 0, D("8.70"), D("0.00")),
                               ("OTA B", 44, 6, 3, D("13.64"), D("7.89")), ("Kurumsal", 21, 1, 0, D("4.76"), D("0.00")), ("Telefon / walk-in", 13, 1, 0, D("7.69"), D("0.00"))],
                              ("T", 260, 43, 8, D("16.54"), D("3.69"))))
            self.assertEqual([(u["tur"], u["kim"]) for u in x["uyarilar"]][:3], [("No-show oranı yüksek", "Kanal: OTA A"), ("No-show oranı yüksek", "Kanal: OTA B"), ("İptal oranı yüksek", "Kanal: OTA A")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Kırılımlar", "İptal Zamanı", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "r.csv"
            y.write_text("Rez No;Rez Tarihi;Check In;Check Out;Statü;İptal Tarihi;Tutar\n1;01.09.2026;01.10.2026;04.10.2026;Konakladı;;1.500\n1;;01.10.2026;;İptal;;\n2;;x;;İptal;;\n"
                         "3;;01.10.2026;;belki;;\n4;05.10.2026;01.10.2026;;İptal;10.10.2026;\n;;;;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["gece"], x["lead"], x["iptal"], x["tutar"]) for x in l], [u["tur"] for u in uy]),
                             ([("1", 3, 30, None, D(1500)), ("4", 1, None, None, None)], ["Rezervasyon iki kez yazılmış", "Satır okunamadı", "Satır okunamadı", "Rezervasyon girişten sonra"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--rezervasyonlar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--min", "1", "--iptal", "10", "--noshow", "2"]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--min", "0"], ["--iptal", "x"], ["--noshow", "101"])], [2] * 3)
            self.assertEqual(main.main(["--rezervasyonlar", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Rez No;Tutar\n1;5\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
