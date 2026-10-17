"""
İptal ve No-Show Analizi — Gerisi Bende kod bloğu
Ön Büro › Ön Büro Müdürü
Rezervasyon iptal ve gelmeme (no-show) oranlarını kanal, oda tipi ve rezervasyon zamanlamasına göre analiz eder.
  İptal oranı %      = iptal edilen rezervasyon ÷ toplam rezervasyon × 100
  No-show oranı %    = gelmeyen ÷ iptal edilmemiş (gelmesi beklenen) rezervasyon × 100
  Oda-gece           = çıkış − giriş (ya da Gece sütunu); kayıp oda-gece ve kayıp gelir iptal + no-show rezervasyonlardan
  Rezervasyon öncesi = giriş − rezervasyon tarihi: 0–7 · 8–30 · 31–90 · 90+ gün
  İptal zamanı       = giriş − iptal tarihi: aynı gün · 1–3 · 4–7 · 8–30 · 30+ gün
  Uyarı              = en az --min (10) rezervasyonlu kanal / oda tipinde iptal oranı --iptal (%25) ya da no-show oranı --noshow (%5) üstünde
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"no": ("rezervasyon no", "rez no", "voucher", "rezervasyon"), "rez": ("rezervasyon tarihi", "rez tarihi", "kayit tarihi", "olusturma tarihi"), "giris": ("giris tarihi", "giris", "check in", "gelis"),
         "cikis": ("cikis tarihi", "cikis", "check out", "ayrilis"), "gece": ("gece", "gece sayisi", "konaklama gecesi"), "kanal": ("kanal", "satis kanali", "acente", "kaynak"),
         "oda": ("oda tipi", "oda", "oda turu"), "durum": ("durum", "rezervasyon durumu", "statu"), "iptal": ("iptal tarihi",), "tutar": ("tutar", "toplam tutar", "rezervasyon tutari", "gelir")}
REZ_BANT = (("0–7 gün", 7), ("8–30 gün", 30), ("31–90 gün", 90), ("90+ gün", None))
IPTAL_BANT = (("Aynı gün", 0), ("1–3 gün", 3), ("4–7 gün", 7), ("8–30 gün", 30), ("30+ gün", None))
SIFIR = Decimal(0)


def _durum(x) -> str | None:
    k = katla(x)
    if not k:
        return None
    if "no show" in k or k in ("noshow", "gelmedi", "gelmeyen", "ns"):
        return "No-show"
    if "iptal" in k or k in ("cancel", "cancelled", "canceled", "cxl"):
        return "İptal"
    if k in ("konakladi", "gerceklesti", "geldi", "check in", "check out", "cikis yapti", "giris yapti", "kesin", "onayli", "aktif", "checked out", "in house"):
        return "Gerçekleşti"
    return None


def _bant(g: int, bantlar) -> str:
    return next(ad for ad, ust in bantlar if ust is None or g <= ust)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("giris", "durum")):
        no = metin(r["no"])
        if not no and not metin(r["giris"]):
            continue
        g, durum = tarih(r["giris"]), _durum(r["durum"])
        rz, c, ip = (tarih(r[a]) if metin(r[a]) else None for a in ("rez", "cikis", "iptal"))
        gece, tutar = (para(r["gece"]) if metin(r["gece"]) else None), (para(r["tutar"]) if metin(r["tutar"]) else None)
        kim = no or f"Satır {r['_satir']}"
        bozuk = any(metin(r[a]) and v is None for a, v in (("rez", rz), ("cikis", c), ("iptal", ip), ("gece", gece), ("tutar", tutar)))
        if g is None or durum is None or bozuk:
            uy.append(uyari("Yüksek", "Satır okunamadı", kim, f"Satır {r['_satir']}: giriş tarihi, durum ('{metin(r['durum'])}') ya da tarih / sayılardan biri okunamadı. Rezervasyon alınmadı."))
            continue
        if no and katla(no) in gorulen:
            uy.append(uyari("Orta", "Rezervasyon iki kez yazılmış", kim, f"Satır {r['_satir']}: ikinci satır alınmadı."))
            continue
        if no:
            gorulen.add(katla(no))
        gece = int(gece) if gece is not None else (c - g).days if c and c > g else 1
        if rz and rz > g:
            uy.append(uyari("Orta", "Rezervasyon girişten sonra", kim, f"Rezervasyon {rz:%d.%m.%Y}, giriş {g:%d.%m.%Y}; rezervasyon öncesi süreye girmedi."))
            rz = None
        if durum == "İptal" and ip and ip > g:
            ip = None
        l.append({"no": no, "rez": rz, "giris": g, "gece": max(gece, 1), "kanal": metin(r["kanal"]) or "Belirtilmemiş", "oda": metin(r["oda"]) or "Belirtilmemiş", "durum": durum, "iptal": ip,
                  "tutar": tutar, "lead": (g - rz).days if rz else None})
    return l, uy


def _oran(l: list[dict]) -> dict:
    n = len(l)
    ip = [x for x in l if x["durum"] == "İptal"]
    ns = [x for x in l if x["durum"] == "No-show"]
    beklenen = n - len(ip)
    kayip = ip + ns
    return {"rez": n, "iptal": len(ip), "noshow": len(ns), "gerceklesen": n - len(kayip), "iptal_oran": yuvarla(Decimal(len(ip)) * 100 / n) if n else None,
            "noshow_oran": yuvarla(Decimal(len(ns)) * 100 / beklenen) if beklenen else None, "gece": sum(x["gece"] for x in l), "kayip_gece": sum(x["gece"] for x in kayip),
            "kayip_gelir": sum((x["tutar"] for x in kayip if x["tutar"] is not None), SIFIR), "lead": yuvarla(Decimal(sum(x["lead"] for x in l if x["lead"] is not None)) / sum(1 for x in l if x["lead"] is not None))
            if any(x["lead"] is not None for x in l) else None}


def analiz(rez: list[dict], en_az: int = 10, iptal: Decimal = Decimal(25), noshow: Decimal = Decimal(5)) -> dict:
    if not rez:
        raise ValueError("Rezervasyon kaydı yok.")
    uy = []

    def kirilim(fn, sira=None):
        g = {}
        for x in rez:
            g.setdefault(fn(x), []).append(x)
        satir = [{"ad": k, **_oran(v)} for k, v in g.items() if k is not None]
        if sira:
            return sorted(satir, key=lambda s: sira.index(s["ad"]))
        return sorted(satir, key=lambda s: (-s["rez"], katla(s["ad"])))
    kanal = kirilim(lambda x: x["kanal"])
    oda = kirilim(lambda x: x["oda"])
    lead = kirilim(lambda x: _bant(x["lead"], REZ_BANT) if x["lead"] is not None else None, [a for a, _ in REZ_BANT])
    ay = sorted(kirilim(lambda x: f"{x['giris']:%Y-%m}"), key=lambda s: s["ad"])
    iptaller = [x for x in rez if x["durum"] == "İptal" and x["iptal"]]
    iptal_zaman = []
    for ad, _ in IPTAL_BANT:
        l = [x for x in iptaller if _bant((x["giris"] - x["iptal"]).days, IPTAL_BANT) == ad]
        iptal_zaman.append({"ad": ad, "adet": len(l), "pay": yuvarla(Decimal(len(l)) * 100 / len(iptaller)) if iptaller else None, "gece": sum(x["gece"] for x in l),
                            "gelir": sum((x["tutar"] for x in l if x["tutar"] is not None), SIFIR)})
    for ad, satirlar in (("Kanal", kanal), ("Oda tipi", oda)):
        for s in satirlar:
            if s["rez"] < en_az:
                continue
            if s["iptal_oran"] is not None and s["iptal_oran"] > iptal:
                uy.append(uyari("Orta", "İptal oranı yüksek", f"{ad}: {s['ad']}", f"{s['rez']} rezervasyondan {s['iptal']} iptal (%{s['iptal_oran']}; eşik %{iptal})."))
            if s["noshow_oran"] is not None and s["noshow_oran"] > noshow:
                uy.append(uyari("Orta", "No-show oranı yüksek", f"{ad}: {s['ad']}", f"Gelmesi beklenen {s['rez'] - s['iptal']} rezervasyondan {s['noshow']} gelmedi (%{s['noshow_oran']}; eşik %{noshow})."))
    son_dakika = sum(x["adet"] for x in iptal_zaman if x["ad"] in ("Aynı gün", "1–3 gün"))
    if iptaller and son_dakika:
        uy.append(uyari("Bilgi", "Son dakika iptalleri", f"{son_dakika} iptal", f"İptallerin %{yuvarla(Decimal(son_dakika) * 100 / len(iptaller))} kadarı girişe 3 gün ya da daha az kala yapıldı; iptal politikası ve overbooking payı için değerlendirin."))
    tarihsiz = sum(1 for x in rez if x["durum"] == "İptal" and not x["iptal"])
    if tarihsiz:
        uy.append(uyari("Bilgi", "İptal tarihi yok", f"{tarihsiz} iptal", "İptal tarihi yazılmamış (ya da girişten sonra) iptaller iptal zamanı tablosuna girmedi."))
    return {"toplam": _oran(rez), "kanal": kanal, "oda": oda, "lead": lead, "ay": ay, "iptal_zaman": iptal_zaman, "uyarilar": uy}


BASLIK = ["Rezervasyon", "İptal", "No-show", "Gerçekleşen", "İptal %", "No-show %", "Oda-Gece", "Kayıp Oda-Gece", "Kayıp Gelir", "Ort. Rezervasyon Öncesi (gün)"]


def rapor_yaz(cikti: Path, r: dict, iptal: Decimal, noshow: Decimal) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kırılımlar"
    _baslik(ws, ["Kırılım", "Ad"] + BASLIK, (16, 22, 10, 7, 8, 10, 8, 9, 9, 10, 12, 12), "C2")

    def yaz(ad, s, kalin=False):
        k = _satir(ws, [ad, s["ad"], s["rez"], s["iptal"], s["noshow"], s["gerceklesen"], s["iptal_oran"], s["noshow_oran"], s["gece"], s["kayip_gece"], s["kayip_gelir"] or None, s["lead"]],
                   para_s=(7, 8, 11, 12), dolgu={7: KIRMIZI if s["iptal_oran"] is not None and s["iptal_oran"] > iptal else None,
                                                 8: KIRMIZI if s["noshow_oran"] is not None and s["noshow_oran"] > noshow else None})
        if kalin:
            for h in ws[k]:
                h.font = Font(bold=True)
    yaz("TOPLAM", {"ad": "", **r["toplam"]}, True)
    for ad, l in (("Kanal", r["kanal"]), ("Oda tipi", r["oda"]), ("Rezervasyon öncesi", r["lead"]), ("Giriş ayı", r["ay"])):
        for s in l:
            yaz(ad, s)
    ws.append([])
    ws.append(["No-show % = gelmeyen ÷ iptal edilmemiş rezervasyon. Kayıp gelir, tutarı yazılmış iptal ve no-show rezervasyonların toplamıdır (tahsil edilen iptal / no-show ücreti düşülmez)."])
    ws = wb.create_sheet("İptal Zamanı")
    _baslik(ws, ["Girişe Kalan Süre", "İptal", "Pay %", "Oda-Gece", "Tutar"], (18, 8, 8, 10, 14))
    for s in r["iptal_zaman"]:
        _satir(ws, [s["ad"], s["adet"], s["pay"], s["gece"], s["gelir"] or None], para_s=(3, 5))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, en_az: int = 10, iptal: Decimal = Decimal(25), noshow: Decimal = Decimal(5)) -> dict:
    rez, uy = oku(yol)
    r = analiz(rez, en_az, iptal, noshow)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, iptal, noshow)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Rezervasyon iptal ve no-show oranlarını kanal, oda tipi ve rezervasyon zamanlamasına göre analiz eder.")
    p.add_argument("--rezervasyonlar", type=Path, help="Rezervasyon No, Rezervasyon Tarihi, Giriş, Çıkış (ya da Gece), Kanal, Oda Tipi, Durum (Konakladı / İptal / No-show), İptal Tarihi, Tutar")
    p.add_argument("--min", type=int, default=10, help="Uyarı için kanal / oda tipinde en az rezervasyon (varsayılan 10)")
    p.add_argument("--iptal", default="25", help="İptal oranı bu yüzdenin üstündeyse uyarılır (varsayılan 25)")
    p.add_argument("--noshow", default="5", help="No-show oranı bu yüzdenin üstündeyse uyarılır (varsayılan 5)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "iptal_noshow_analizi.xlsx")
    a = p.parse_args(argv)
    yol = a.rezervasyonlar or ORNEK / "rezervasyonlar.csv"
    ip, ns = para(a.iptal), para(a.noshow)
    if a.min < 1 or ip is None or not 0 <= ip <= 100 or ns is None or not 0 <= ns <= 100:
        print("[X] --min en az 1; --iptal ve --noshow 0-100 arası olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, a.min, ip, ns)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {t['rez']} rezervasyon · iptal %{t['iptal_oran']} · no-show %{t['noshow_oran']} · kayıp oda-gece {t['kayip_gece']} · kayıp gelir {t['kayip_gelir']:,.2f}")
    for k in r["kanal"]:
        print(f"     {k['ad']}: {k['rez']} rez · iptal %{k['iptal_oran']} · no-show %{k['noshow_oran']}")
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
