import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
H41 = date(2026, 10, 5)


def h(n):
    return H41 + timedelta(weeks=n - 41)


def ie(emir, hafta, yuk, merkez="CNC", termin=None):
    return {"emir": emir, "merkez": merkez, "anahtar": main.katla(merkez), "hafta": h(hafta), "yuk": D(yuk), "termin": h(termin) if termin else None, "ilk_hafta": h(hafta)}


def kap(**kw):
    return {main.katla(k): {"ad": k, "genel": D(v), "ozel": {}} for k, v in kw.items()}


def oneri(r):
    return [(o["emir"], main.hafta_ad(o["kaynak"])[-3:], main.hafta_ad(o["hedef"])[-3:], o["yon"], o["risk"]) for o in r["oneriler"]]


class AnalizTesti(unittest.TestCase):
    def test_hafta(self):
        self.assertEqual([main.hafta(x) for x in ("2026-H41", "2026-W41", "2026/41", "H41-2026", "41/2026", "07.10.2026", date(2026, 10, 11), "2026-H60", "ekim")],
                         [H41, H41, H41, H41, H41, H41, H41, None, None])
        self.assertEqual(main.hafta_ad(H41), "2026-H41")

    def test_yukleme_ve_durum(self):
        r = main.analiz([ie("A", 41, 30), ie("B", 41, 25), ie("C", 42, 72), ie("D", 43, 80)], kap(CNC=80))
        self.assertEqual([(t["yuk"], t["oran"], t["durum"]) for t in r["tablo"]], [(D(55), D("68.75"), "Uygun"), (D(72), D("90.00"), "Sınırda"), (D(80), D("100.00"), "Sınırda")])
        self.assertEqual((r["darbogaz"], r["oneriler"]), ([], []))                                   # tam kapasite darboğaz değildir

    def test_one_cekme_en_kucuk_yeterli_emir(self):
        r = main.analiz([ie("A", 41, 30), ie("B", 41, 25), ie("C", 42, 40), ie("D", 42, 35), ie("E", 42, 20), ie("F", 43, 70)], kap(CNC=80))
        self.assertEqual(oneri(r), [("E", "H42", "H41", "Öne çekme", False)])                         # aşım 15: tek başına kapatan en küçük emir 20 saatlik E
        t = r["tablo"]
        self.assertEqual([(x["oran"], x["asim"], x["sonra_oran"], x["sonra_durum"]) for x in t[:2]], [(D("68.75"), D(0), D("93.75"), "Sınırda"), (D("118.75"), D(15), D("93.75"), "Sınırda")])

    def test_termin_riski_olmayan_tercih_edilir(self):
        e = [ie("X", 42, 100, "Montaj"), ie("A", 43, 70, "Montaj", 45), ie("B", 43, 60, "Montaj", 43), ie("Y", 44, 30, "Montaj")]
        self.assertEqual(oneri(main.analiz(e, kap(Montaj=120))), [("A", "H43", "H44", "Erteleme", False)])
        e = [ie("X", 42, 100, "Montaj"), ie("A", 43, 70, "Montaj", 43), ie("B", 43, 60, "Montaj", 43), ie("Y", 44, 30, "Montaj")]
        self.assertEqual(oneri(main.analiz(e, kap(Montaj=120))), [("B", "H43", "H44", "Erteleme", True)])

    def test_cozulemeyen_ve_coklu_kaydirma(self):
        r = main.analiz([ie("A", 41, 45, "Pres"), ie("B", 42, 38, "Pres"), ie("C", 43, 10, "Pres")], kap(Pres=40))
        self.assertEqual((oneri(r), [(c["merkez"], c["asim"]) for c in r["cozulemeyen"]]), ([], [("Pres", D(5))]))
        r = main.analiz([ie("A", 41, 10), ie("B", 42, 30), ie("C", 42, 30), ie("D", 42, 30), ie("E", 42, 30), ie("F", 43, 20)], kap(CNC=80))
        self.assertEqual(sorted(o[2] for o in oneri(r)), ["H41", "H41"])                              # aşım 40: iki emir de boş olan önceki haftaya
        self.assertEqual([t["sonra"] for t in r["tablo"]], [D(70), D(60), D(20)])

    def test_ozel_kapasite_ve_tanimsiz_merkez(self):
        k = kap(CNC=80)
        k["cnc"]["ozel"][h(42)] = D(40)
        r = main.analiz([ie("A", 41, 10), ie("B", 42, 50), ie("C", 41, 5, "Kaynak")], k)
        self.assertEqual(([(t["kapasite"], t["durum"]) for t in r["tablo"]], oneri(r), [u["tur"] for u in r["uyarilar"]]),
                         ([(D(80), "Uygun"), (D(40), "Darboğaz")], [("B", "H42", "H41", "Öne çekme", False)], ["İş merkezinin kapasitesi tanımsız"]))
        with self.assertRaises(ValueError):
            main.analiz([ie("C", 41, 5, "Kaynak")], k)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "k.xlsx"
            r = main.calistir(ORNEK / "is_emirleri.csv", ORNEK / "kapasiteler.csv", cikti)
            self.assertEqual([(t["merkez"], main.hafta_ad(t["hafta"]), t["oran"], t["asim"]) for t in r["darbogaz"]],
                             [("CNC", "2026-H42", D("118.75"), D(15)), ("Pres", "2026-H41", D("112.50"), D(5)), ("Montaj", "2026-H43", D("108.33"), D(10))])
            self.assertEqual(oneri(r), [("IE-05", "H42", "H41", "Öne çekme", False), ("IE-13", "H43", "H41", "Öne çekme", False)])               # H42 dolu, H44 termini aşar; boş olan H41'e
            self.assertEqual(([(c["merkez"], c["asim"]) for c in r["cozulemeyen"]], len(r["ufuk"]), r["uyarilar"]), ([("Pres", D(5))], 4, []))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Yükleme Tablosu", "Darboğazlar", "Kaydırma Önerileri", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, k = Path(tmp) / "y.csv", Path(tmp) / "k.csv"
            y.write_text("İş Emri;İş Merkezi;Hafta;Yük (saat);Adet;Birim Süre (dk);Hazırlık (dk);Termin Haftası\nA;CNC;2026-H41;;600;3;30;yakında\nB;CNC;ekim;10;;;;\nC;CNC;2026-H41;çok;;;;\n"
                         "D;CNC;06.10.2026;12,5;;;;2026-H42\n", encoding="utf-8")
            k.write_text("İş Merkezi;Haftalık Kapasite (saat);Hafta\nCNC;80;\nCNC;40;2026-H42\nPres;yok;\nPres;10;bozuk\n", encoding="utf-8")
            e, kp, uy = main.oku(y, k)
            self.assertEqual([(x["emir"], x["hafta"], x["yuk"], x["termin"]) for x in e], [("A", H41, D("30.50"), None), ("D", H41, D("12.50"), h(42))])   # (600 × 3 + 30) ÷ 60
            self.assertEqual((kp["cnc"]["genel"], kp["cnc"]["ozel"], "pres" in kp), (D(80), {h(42): D(40)}, False))
            self.assertEqual(sorted(u["tur"] for u in uy), ["Kapasite satırı okunamadı", "Kapasite satırı okunamadı", "Termin haftası okunamadı", "İş emri satırı okunamadı", "İş emri satırı okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--is-emirleri", str(y), "--kapasiteler", str(k), "--cikti", c]
            self.assertEqual(main.main(arg + ["--uyari", "50"]), 0)
            self.assertEqual(main.main(arg + ["--uyari", "0"]), 2)
            self.assertEqual(main.main(["--is-emirleri", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("İş Merkezi;Hafta;Yük (saat)\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
