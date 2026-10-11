"""
Kapasite Darboğaz Analizi — Workers / Workless kod bloğu
Üretim Planlama › Üretim Planlama Uzmanı

İş merkezi × hafta bazında yükleme oranını hesaplar; darboğaz haftalarını ve iş emri kaydırma önerilerini çıkarır.
  Yük (saat)   = "Yük (saat)" sütunu;  yoksa (adet × birim süre dk + hazırlık dk) ÷ 60
  Yükleme %    = haftalık yük ÷ haftalık kapasite × 100;   > 100 "Darboğaz",  ≥ --uyari (90) "Sınırda"
  Kaydırma     = darboğaz haftasındaki bir iş emrini, tamamı sığan en yakın haftaya taşır (eşit uzaklıkta önce öne çekme).
                 Aşımı tek başına kapatan en küçük iş emri önce denenir; termin haftasını aşmayan seçenek tercih edilir.
  Çözülemeyen  = hiçbir haftaya sığmayan aşım: fazla mesai, fason veya iş emrini bölme gerekir
Kod yalnız kapasiteye bakar: operasyon sırası (rota), malzeme, kalıp ve personel uygunluğu kontrol edilmez. Öneriler taslaktır.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


YUK_SUTUN = {"emir": ("is emri", "is emri no", "siparis", "siparis no"), "merkez": ("is merkezi", "makine", "hat", "kaynak"), "hafta": ("hafta", "plan haftasi", "tarih"),
             "yuk": ("yuk saat", "yuk", "sure saat", "saat"), "adet": ("adet", "miktar"), "birim": ("birim sure dk", "birim sure", "cevrim suresi dk"),
             "hazirlik": ("hazirlik dk", "hazirlik", "ayar suresi dk"), "termin": ("termin haftasi", "termin", "en gec hafta")}
KAP_SUTUN = {"merkez": ("is merkezi", "makine", "hat", "kaynak"), "kapasite": ("haftalik kapasite saat", "haftalik kapasite", "kapasite saat", "kapasite"), "hafta": ("hafta",)}
Z = Decimal(0)


def hafta(x) -> date | None:
    """Haftayı o haftanın pazartesi tarihine çevirir: '2026-H42', '2026-W42', '42/2026' veya bir tarih."""
    if isinstance(x, (date, datetime)):
        t = x.date() if isinstance(x, datetime) else x
        return t - timedelta(days=t.weekday())
    s = metin(x)
    t = tarih(s) if re.fullmatch(r"\d{1,4}[./-]\d{1,2}[./-]\d{2,4}", s) else None
    if t:
        return t - timedelta(days=t.weekday())
    m = re.fullmatch(r"(\d{4})\D{1,3}(\d{1,2})", s) or None
    if m:
        y, h = int(m[1]), int(m[2])
    else:
        m = re.fullmatch(r"\D{0,2}(\d{1,2})\D{1,3}(\d{4})", s)
        if not m:
            return None
        y, h = int(m[2]), int(m[1])
    try:
        return date.fromisocalendar(y, h, 1)
    except ValueError:
        return None


def hafta_ad(t: date) -> str:
    y, h, _ = t.isocalendar()
    return f"{y}-H{h:02d}"


def oku(yuk_yol: Path, kap_yol: Path) -> tuple[list[dict], dict, list[dict]]:
    uy, emirler, kap = [], [], {}
    for r in kayitlar(kap_yol, KAP_SUTUN, ("merkez", "kapasite")):
        m = metin(r["merkez"])
        if not m:
            continue
        k, h = para(r["kapasite"]), hafta(r["hafta"])
        if k is None or k < 0 or (metin(r["hafta"]) and not h):
            uy.append(uyari("Yüksek", "Kapasite satırı okunamadı", m, f"{kap_yol.name} satır {r['_satir']}: kapasite '{metin(r['kapasite'])}', hafta '{metin(r['hafta'])}'. Satır kullanılmadı."))
            continue
        kap.setdefault(katla(m), {"ad": m, "genel": None, "ozel": {}})
        if h:
            kap[katla(m)]["ozel"][h] = k
        else:
            kap[katla(m)]["genel"] = k
    for r in kayitlar(yuk_yol, YUK_SUTUN, ("merkez", "hafta")):
        m = metin(r["merkez"])
        if not m:
            continue
        kim = f"{metin(r['emir']) or m} (satır {r['_satir']})"
        h, y = hafta(r["hafta"]), para(r["yuk"])
        if y is None and para(r["adet"]) is not None and para(r["birim"]) is not None:
            y = (para(r["adet"]) * para(r["birim"]) + (para(r["hazirlik"]) or Z)) / 60
        if not h or y is None or y < 0:
            uy.append(uyari("Yüksek", "İş emri satırı okunamadı", kim, f"Hafta '{metin(r['hafta'])}', yük '{metin(r['yuk'])}'. Hafta 2026-H42 gibi veya tarih; yük saat (ya da adet × birim süre) olmalı."))
            continue
        t = hafta(r["termin"])
        if metin(r["termin"]) and not t:
            uy.append(uyari("Orta", "Termin haftası okunamadı", kim, f"'{metin(r['termin'])}'. Termin kontrolü yapılmadı."))
        emirler.append({"emir": metin(r["emir"]) or f"satır {r['_satir']}", "merkez": m, "anahtar": katla(m), "hafta": h, "yuk": yuvarla(y), "termin": t, "ilk_hafta": h})
    return emirler, kap, uy


def analiz(emirler: list[dict], kap: dict, uyari_esik: Decimal = Decimal(90)) -> dict:
    uy = []
    for a in sorted({e["anahtar"] for e in emirler} - {k for k, v in kap.items() if v["genel"] is not None or v["ozel"]}):
        ad = next(e["merkez"] for e in emirler if e["anahtar"] == a)
        uy.append(uyari("Yüksek", "İş merkezinin kapasitesi tanımsız", ad, "Kapasite dosyasında bu iş merkezi yok; yükleri analize alınmadı."))
    emirler = [e for e in emirler if e["anahtar"] in kap and (kap[e["anahtar"]]["genel"] is not None or kap[e["anahtar"]]["ozel"])]
    if not emirler:
        raise ValueError("Kapasitesi tanımlı iş merkezine ait iş emri yok.")
    ilk, son = min(e["hafta"] for e in emirler), max(e["hafta"] for e in emirler)
    ufuk = [ilk + timedelta(weeks=i) for i in range((son - ilk).days // 7 + 1)]
    merkezler = list(dict.fromkeys(e["anahtar"] for e in emirler))

    def kapasite(m, h):
        return kap[m]["ozel"].get(h, kap[m]["genel"] if kap[m]["genel"] is not None else Z)

    def yuk(m, h):
        return sum((e["yuk"] for e in emirler if e["anahtar"] == m and e["hafta"] == h), Z)

    def durum(y, k):
        return "Darboğaz" if y > k else "Sınırda" if k and y / k * 100 >= uyari_esik else "Uygun"

    once = {(m, h): yuk(m, h) for m in merkezler for h in ufuk}
    oneriler, cozulemeyen = [], []
    for m in merkezler:
        for i, h in enumerate(ufuk):
            while yuk(m, h) > kapasite(m, h):
                asim = yuk(m, h) - kapasite(m, h)
                buradaki = [e for e in emirler if e["anahtar"] == m and e["hafta"] == h and e["ilk_hafta"] == h]
                adaylar = sorted((e for e in buradaki if e["yuk"] >= asim), key=lambda e: e["yuk"]) + sorted((e for e in buradaki if e["yuk"] < asim), key=lambda e: -e["yuk"])
                secim = None
                for e in adaylar:
                    riskli = None
                    for uzak in range(1, len(ufuk)):
                        for j in (i - uzak, i + uzak):
                            if 0 <= j < len(ufuk) and kapasite(m, ufuk[j]) - yuk(m, ufuk[j]) >= e["yuk"]:
                                risk = e["termin"] is not None and ufuk[j] > e["termin"]
                                if not risk:
                                    secim = (e, ufuk[j], False)
                                    break
                                riskli = riskli or (e, ufuk[j], True)
                        if secim and secim[0] is e:
                            break
                    if secim and not secim[2]:
                        break
                    secim = secim or riskli
                if not secim:
                    cozulemeyen.append({"merkez": kap[m]["ad"], "hafta": h, "asim": asim})
                    break
                e, hedef, risk = secim
                oneriler.append({"emir": e["emir"], "merkez": kap[m]["ad"], "kaynak": h, "hedef": hedef, "yuk": e["yuk"], "yon": "Öne çekme" if hedef < h else "Erteleme", "risk": risk,
                                 "termin": e["termin"]})
                e["hafta"] = hedef
    tablo = []
    for m in merkezler:
        for h in ufuk:
            k, y0, y1 = kapasite(m, h), once[(m, h)], yuk(m, h)
            tablo.append({"merkez": kap[m]["ad"], "hafta": h, "kapasite": k, "yuk": y0, "oran": yuvarla(y0 / k * 100) if k else None, "durum": durum(y0, k), "asim": max(y0 - k, Z),
                          "sonra": y1, "sonra_oran": yuvarla(y1 / k * 100) if k else None, "sonra_durum": durum(y1, k)})
    return {"tablo": tablo, "ufuk": ufuk, "oneriler": oneriler, "cozulemeyen": cozulemeyen, "darbogaz": [t for t in tablo if t["durum"] == "Darboğaz"], "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    renk = {"Darboğaz": KIRMIZI, "Sınırda": SARI, "Uygun": None}
    ws = wb.active
    ws.title = "Yükleme Tablosu"
    _baslik(ws, ["İş Merkezi", "Hafta", "Hafta Başı", "Kapasite (saat)", "Yük (saat)", "Yükleme %", "Durum", "Aşım (saat)", "Öneri Sonrası Yük", "Öneri Sonrası %", "Öneri Sonrası Durum"],
            (22, 11, 12, 13, 12, 11, 12, 12, 14, 12, 16))
    for t in r["tablo"]:
        _satir(ws, [t["merkez"], hafta_ad(t["hafta"]), t["hafta"], t["kapasite"], t["yuk"], t["oran"], t["durum"], t["asim"] or None, t["sonra"], t["sonra_oran"], t["sonra_durum"]],
               para_s=(4, 5, 6, 8, 9, 10), tarih_s=(3,), dolgu={7: renk[t["durum"]], 11: renk[t["sonra_durum"]]})
    ws = wb.create_sheet("Darboğazlar")
    _baslik(ws, ["İş Merkezi", "Hafta", "Kapasite (saat)", "Yük (saat)", "Yükleme %", "Aşım (saat)", "Öneri Sonrası %", "Sonuç"], (22, 11, 13, 12, 11, 12, 12, 70))
    kalan = {(c["merkez"], c["hafta"]): c["asim"] for c in r["cozulemeyen"]}
    for t in r["darbogaz"]:
        k = kalan.get((t["merkez"], t["hafta"]))
        _satir(ws, [t["merkez"], hafta_ad(t["hafta"]), t["kapasite"], t["yuk"], t["oran"], t["asim"], t["sonra_oran"],
                    f"Kaydırmayla çözülemedi: {k:f} saat için fazla mesai, fason veya iş emrini bölme gerekir." if k else "Kaydırma önerisiyle kapasite içine alındı."],
               para_s=(3, 4, 5, 6, 7), dolgu={8: KIRMIZI if k else YESIL}, sar=(8,))
    ws = wb.create_sheet("Kaydırma Önerileri")
    _baslik(ws, ["İş Emri", "İş Merkezi", "Mevcut Hafta", "Önerilen Hafta", "Yük (saat)", "Yön", "Termin Haftası", "Termin Riski", "Karar (siz doldurun)"], (16, 22, 12, 13, 11, 12, 13, 12, 26))
    for o in r["oneriler"]:
        _satir(ws, [o["emir"], o["merkez"], hafta_ad(o["kaynak"]), hafta_ad(o["hedef"]), o["yuk"], o["yon"], hafta_ad(o["termin"]) if o["termin"] else "", "Var" if o["risk"] else "", ""],
               para_s=(5,), dolgu={8: KIRMIZI if o["risk"] else None})
    ws.append([])
    ws.append(["Öneriler yalnız kapasiteye göredir. Rota sırası, malzeme, kalıp ve personel uygunluğunu kontrol etmeden uygulamayın."])
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yuk_yol: Path, kap_yol: Path, cikti: Path, uyari_esik: Decimal = Decimal(90)) -> dict:
    emirler, kap, uy = oku(yuk_yol, kap_yol)
    if not emirler:
        raise ValueError(f"{yuk_yol.name}: okunabilen iş emri yok.")
    r = analiz(emirler, kap, uyari_esik)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, r["uyarilar"])
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="İş merkezi bazında haftalık yükleme oranını, darboğazları ve kaydırma önerilerini çıkarır.")
    p.add_argument("--is-emirleri", type=Path, default=ORNEK / "is_emirleri.csv", help="İş Emri, İş Merkezi, Hafta, Yük (saat) veya Adet + Birim Süre (dk) + Hazırlık (dk), Termin Haftası")
    p.add_argument("--kapasiteler", type=Path, default=ORNEK / "kapasiteler.csv", help="İş Merkezi, Haftalık Kapasite (saat), Hafta (yalnız o haftaya özel kapasite için)")
    p.add_argument("--uyari", default="90", help="Sınırda sayılacak yükleme %% (varsayılan 90)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kapasite_darbogaz.xlsx")
    a = p.parse_args(argv)
    for y in (a.is_emirleri, a.kapasiteler):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    esik = para(a.uyari)
    if esik is None or not 0 < esik <= 100:
        print("[X] --uyari 0 ile 100 arasında olmalı")
        return 2
    try:
        r = calistir(a.is_emirleri, a.kapasiteler, a.cikti, esik)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(r['ufuk'])} hafta · {len(r['darbogaz'])} darboğaz · {len(r['oneriler'])} kaydırma önerisi · {len(r['cozulemeyen'])} çözülemeyen")
    for t in r["darbogaz"]:
        print(f"[!] {t['merkez']} · {hafta_ad(t['hafta'])}: yükleme %{t['oran']} (aşım {t['asim']:f} saat) → öneri sonrası %{t['sonra_oran']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
