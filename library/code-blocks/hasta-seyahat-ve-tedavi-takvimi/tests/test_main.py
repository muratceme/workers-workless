import sys
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def e(hasta, tur, bas, bit=None, saatli=True):
    b = datetime(2026, 10, *bas)
    return {"hasta": hasta, "tur": tur, "ham_tur": tur, "bas": b, "saatli": saatli, "bit": datetime(2026, 10, *bit) if bit else None, "bit_saatli": False, "yer": "", "aciklama": ""}


def veri():
    return [e("H1", "Varış uçuşu", (10, 8, 0)), e("H1", "Transfer", (10, 8, 30)), e("H1", "Konaklama", (10,), (11,), False), e("H1", "Muayene", (10, 14, 0)),
            e("H1", "Yatış", (11,), (12,), False), e("H1", "Operasyon", (11, 9, 0)), e("H1", "Konaklama", (13,), (15,), False), e("H1", "Kontrol", (14, 10, 0)), e("h1", "Dönüş uçuşu", (15, 18, 0)),
            e("H2", "Varış uçuşu", (2, 10, 0)), e("H2", "Muayene", (1, 9, 0)), e("H2", "Tetkik", (3, 10, 0)), e("H2", "Kontrol", (3, 10, 0))]


class KontrolTesti(unittest.TestCase):
    def test_kontroller(self):
        r = main.kontrol(veri(), date(2026, 10, 9), 3, 5)
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Bilgi", "Dönüş uçuşu yok", "H2"), ("Bilgi", "Uçuş günü transfer yok", "H1"), ("Orta", "Aynı anda iki etkinlik", "H2"), ("Orta", "Konaklama boşluğu", "H1"),
                          ("Orta", "Operasyon sonrası uçuş erken", "H1"), ("Yüksek", "Randevu varıştan önce", "H2")])
        self.assertIn("12.10.2026", next(u["aciklama"] for u in r["uyarilar"] if u["tur"] == "Konaklama boşluğu"))
        self.assertEqual(([(x["tur"], x["bas"].day) for x in r["yaklasan"]], r["hasta_sayisi"]),
                         ([("Konaklama", 10), ("Varış uçuşu", 10), ("Transfer", 10), ("Muayene", 10), ("Yatış", 11), ("Operasyon", 11)], 2))
        self.assertNotIn("Operasyon sonrası uçuş erken", [u["tur"] for u in main.kontrol(veri(), date(2026, 10, 9))["uyarilar"]])
        r = main.kontrol([e("H3", "Operasyon", (5, 9, 0)), e("H3", "Kontrol", (20, 9, 0)), e("H3", "Dönüş uçuşu", (19, 9, 0)), e("H3", "Varış uçuşu", (4, 9, 0))], date(2026, 10, 1))
        self.assertEqual(sorted(u["tur"] for u in r["uyarilar"]), ["Konaklama boşluğu", "Operasyon öncesi muayene yok", "Randevu dönüşten sonra", "Uçuş günü transfer yok", "Uçuş günü transfer yok"])
        with self.assertRaises(ValueError):
            main.kontrol([], date(2026, 10, 1))

    def test_yardimcilar(self):
        self.assertEqual([main._tur(x) for x in ("Geliş uçuşu", "Dönüş", "Otel", "Havalimanı transfer (varış)", "Ameliyat", "Kan tahlili", "Pansuman", "Rehber")],
                         ["Varış uçuşu", "Dönüş uçuşu", "Konaklama", "Transfer", "Operasyon", "Tetkik", "Kontrol", "Rehber"])
        self.assertEqual([main.zaman(*x) for x in (("10.10.2026 08:30",), ("10.10.2026", "14.05"), ("10.10.2026",), ("x",), ("10.10.2026", "25:00"))],
                         [(datetime(2026, 10, 10, 8, 30), True), (datetime(2026, 10, 10, 14, 5), True), (datetime(2026, 10, 10), False), (None, False), (datetime(2026, 10, 10), False)])
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "t.ics"
            main.ics_yaz(y, veri()[:3])
            ics = y.read_text(encoding="utf-8")
            self.assertEqual((ics.count("BEGIN:VEVENT"), "DTSTART:20261010T083000" in ics, "DTSTART;VALUE=DATE:20261010" in ics, "DTEND;VALUE=DATE:20261011" in ics), (3, True, True, True))
            self.assertEqual(main._ics_metin("a;b,c\\d"), "a\\;b\\,c\\\\d")


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "hasta_etkinlikleri.csv", cikti, date(2026, 10, 9), 3, 5)
            self.assertEqual((r["hasta_sayisi"], len(r["takvim"]), len(r["yaklasan"])), (3, 24, 14))
            self.assertEqual([(u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Randevu varıştan önce", "UH-2026-052"), ("Aynı anda iki etkinlik", "UH-2026-047"), ("Konaklama boşluğu", "UH-2026-047"),
                              ("Operasyon sonrası uçuş erken", "UH-2026-047"), ("Dönüş uçuşu yok", "UH-2026-052"), ("Uçuş günü transfer yok", "UH-2026-047")])
            self.assertEqual((load_workbook(cikti).sheetnames, cikti.with_suffix(".ics").exists()), (["Takvim", "Yaklaşanlar", "Uyarılar"], True))

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "e.csv"
            y.write_text("Hasta Kodu;Etkinlik;Tarih;Saat;Bitiş\nA;Muayene;10.10.2026;09:30;\nA;Otel;10.10.2026;;12.10.2026\nA;Muayene;x;;\nA;Otel;12.10.2026;;10.10.2026\n;Muayene;10.10.2026;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["tur"], x["bas"], x["saatli"], x["bit"]) for x in l], [u["tur"] for u in uy]),
                             ([("Muayene", datetime(2026, 10, 10, 9, 30), True, None), ("Konaklama", datetime(2026, 10, 10), False, datetime(2026, 10, 12))], ["Satır okunamadı", "Bitiş başlangıçtan önce"]))
            c = Path(tmp) / "x.xlsx"
            arg = ["--etkinlikler", str(y), "--cikti", str(c)]
            self.assertEqual(main.main(arg + ["--tarih", "09.10.2026", "--gun", "2", "--ucus-sure", "3"]), 0)
            self.assertTrue(c.with_suffix(".ics").exists())
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--gun", "-1"], ["--ucus-sure", "-1"])], [2] * 3)
            self.assertEqual(main.main(["--etkinlikler", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Hasta Kodu;Saat\nA;x\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", str(c)]), 0)


if __name__ == "__main__":
    unittest.main()
