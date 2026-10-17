"""
Hasta Seyahat ve Tedavi Takvimi — Gerisi Bende kod bloğu
Uluslararası Hasta Hizmetleri › Uluslararası Hasta Koordinatörü
Uçuş, konaklama, transfer, muayene ve operasyon kayıtlarından hasta bazında takvim, hatırlatma listesi ve .ics takvim dosyası oluşturur.
  Kontroller = tıbbi randevu varıştan önce ya da dönüşten sonra · operasyondan önce muayene yok · aynı anda iki etkinlik ·
               varış-dönüş arasında konaklama / yatışla karşılanmayan gece · uçuş günü transfer yok ·
               operasyon ile dönüş uçuşu arasında --ucus-sure günden az (yalnız değer verilirse; süre hekim kararıdır)
  Hatırlatmalar = rapor tarihinden itibaren --gun (3) gün içindeki etkinlikler
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"hasta": ("hasta", "hasta adi", "hasta kodu", "dosya no", "ad soyad"), "tur": ("tur", "etkinlik", "etkinlik turu", "islem"), "bas": ("baslangic", "tarih", "baslangic tarihi", "giris"),
         "saat": ("saat", "baslangic saati"), "bit": ("bitis", "bitis tarihi", "cikis"), "bit_saat": ("bitis saati",), "yer": ("yer", "adres", "hastane", "otel", "havalimani"),
         "aciklama": ("aciklama", "not", "detay", "ucus no")}
TURLER = (("Transfer", ("transfer", "karsilama")), ("Varış uçuşu", ("varis", "gelis", "arrival")), ("Dönüş uçuşu", ("donus", "departure", "gidis", "kalkis")),
          ("Konaklama", ("konaklama", "otel", "hotel")), ("Yatış", ("yatis", "hastane yatisi")), ("Operasyon", ("operasyon", "ameliyat", "cerrahi", "girisim")),
          ("Muayene", ("muayene", "konsultasyon", "gorusme")), ("Tetkik", ("tetkik", "tahlil", "goruntuleme", "test")), ("Kontrol", ("kontrol", "pansuman", "takip")))
TIBBI = ("Muayene", "Tetkik", "Operasyon", "Kontrol", "Yatış")


def _tur(x) -> str:
    k = katla(x)
    return next((ad for ad, anahtar in TURLER if any(a in k for a in anahtar)), metin(x) or "Diğer")


def zaman(t, s=None) -> tuple[datetime | None, bool]:
    """(tarih-saat, saat yazılmış mı)."""
    g = tarih(t)
    if g is None:
        return None, False
    kaynak = metin(s) if metin(s) else (str(t)[10:] if isinstance(t, str) else (t.strftime("%H:%M") if isinstance(t, datetime) and (t.hour or t.minute) else ""))
    m = re.search(r"(\d{1,2})[:.](\d{2})", kaynak)
    if m and int(m.group(1)) < 24 and int(m.group(2)) < 60:
        return datetime(g.year, g.month, g.day, int(m.group(1)), int(m.group(2))), True
    return datetime(g.year, g.month, g.day), False


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("hasta", "tur", "bas")):
        hasta = metin(r["hasta"])
        if not hasta:
            continue
        bas, saatli = zaman(r["bas"], r["saat"])
        bit, bit_saatli = zaman(r["bit"], r["bit_saat"]) if metin(r["bit"]) else (None, False)
        if bas is None or (metin(r["bit"]) and bit is None):
            uy.append(uyari("Yüksek", "Satır okunamadı", hasta, f"Satır {r['_satir']}: başlangıç ya da bitiş tarihi okunamadı. Etkinlik alınmadı."))
            continue
        if bit and bit < bas:
            uy.append(uyari("Yüksek", "Bitiş başlangıçtan önce", hasta, f"Satır {r['_satir']}: {metin(r['tur'])} bitişi başlangıçtan önce. Etkinlik alınmadı."))
            continue
        l.append({"hasta": hasta, "tur": _tur(r["tur"]), "ham_tur": metin(r["tur"]), "bas": bas, "saatli": saatli, "bit": bit, "bit_saatli": bit_saatli, "yer": metin(r["yer"]), "aciklama": metin(r["aciklama"])})
    return l, uy


def kontrol(etkinlikler: list[dict], gun: date, ileri: int = 3, ucus_sure: int = 0) -> dict:
    if not etkinlikler:
        raise ValueError("Etkinlik kaydı yok.")
    uy, hastalar = [], {}
    for e in etkinlikler:
        hastalar.setdefault(katla(e["hasta"]), []).append(e)
    for k, l in hastalar.items():
        l.sort(key=lambda e: (e["bas"], e["tur"]))
        ad = l[0]["hasta"]
        varis = next((e for e in l if e["tur"] == "Varış uçuşu"), None)
        donus = next((e for e in reversed(l) if e["tur"] == "Dönüş uçuşu"), None)
        if varis is None:
            uy.append(uyari("Bilgi", "Varış uçuşu yok", ad, "Varış kaydı olmadığı için randevuların varıştan önce olup olmadığı kontrol edilemedi."))
        if donus is None:
            uy.append(uyari("Bilgi", "Dönüş uçuşu yok", ad, "Dönüş kaydı yok; konaklama sürekliliği kontrol edilemedi."))
        for e in l:
            if e["tur"] in TIBBI and varis and e["bas"] < varis["bas"]:
                uy.append(uyari("Yüksek", "Randevu varıştan önce", ad, f"{e['tur']} {e['bas']:%d.%m.%Y %H:%M}, varış {varis['bas']:%d.%m.%Y %H:%M}."))
            if e["tur"] in TIBBI and donus and (e["bit"] or e["bas"]) > donus["bas"]:
                uy.append(uyari("Yüksek", "Randevu dönüşten sonra", ad, f"{e['tur']} {e['bas']:%d.%m.%Y %H:%M}, dönüş uçuşu {donus['bas']:%d.%m.%Y %H:%M}."))
        for op in (e for e in l if e["tur"] == "Operasyon"):
            if not any(e["tur"] == "Muayene" and e["bas"] < op["bas"] for e in l):
                uy.append(uyari("Orta", "Operasyon öncesi muayene yok", ad, f"Operasyon {op['bas']:%d.%m.%Y}; öncesinde muayene kaydı yok."))
            if ucus_sure and donus and (donus["bas"].date() - op["bas"].date()).days < ucus_sure:
                uy.append(uyari("Orta", "Operasyon sonrası uçuş erken", ad, f"Operasyon {op['bas']:%d.%m.%Y}, dönüş {donus['bas']:%d.%m.%Y} ({(donus['bas'].date() - op['bas'].date()).days} gün; "
                                                                             f"istenen en az {ucus_sure}). Uçuş uygunluğunu hekimle teyit edin."))
        zamanli = [e for e in l if e["saatli"] and e["tur"] not in ("Konaklama",)]
        for i, a in enumerate(zamanli):
            for b in zamanli[i + 1:]:
                a_bit = a["bit"] if a["bit"] and a["bit_saatli"] else a["bas"]
                if b["bas"] == a["bas"] or (a_bit > a["bas"] and b["bas"] < a_bit and a["tur"] != "Yatış" and b["tur"] != "Yatış"):
                    uy.append(uyari("Orta", "Aynı anda iki etkinlik", ad, f"{a['tur']} {a['bas']:%d.%m.%Y %H:%M} ile {b['tur']} {b['bas']:%d.%m.%Y %H:%M} çakışıyor."))
        if varis and donus:
            geceler = {varis["bas"].date() + timedelta(days=i) for i in range((donus["bas"].date() - varis["bas"].date()).days)}
            kapsanan = set()
            for e in l:
                if e["tur"] in ("Konaklama", "Yatış") and e["bit"]:
                    kapsanan |= {e["bas"].date() + timedelta(days=i) for i in range((e["bit"].date() - e["bas"].date()).days)}
            bos = sorted(geceler - kapsanan)
            if bos:
                uy.append(uyari("Orta", "Konaklama boşluğu", ad, f"Konaklama ya da yatışla karşılanmayan gece: {', '.join(f'{g:%d.%m.%Y}' for g in bos[:10])}" + (" …" if len(bos) > 10 else "") + "."))
            for ucus in (varis, donus):
                if not any(e["tur"] == "Transfer" and e["bas"].date() == ucus["bas"].date() for e in l):
                    uy.append(uyari("Bilgi", "Uçuş günü transfer yok", ad, f"{ucus['tur']} {ucus['bas']:%d.%m.%Y} için transfer kaydı yok."))
    son = datetime.combine(gun + timedelta(days=ileri), datetime.max.time())
    yaklasan = sorted((e for e in etkinlikler if datetime.combine(gun, datetime.min.time()) <= e["bas"] <= son), key=lambda e: (e["bas"], katla(e["hasta"])))
    takvim = sorted(etkinlikler, key=lambda e: (katla(e["hasta"]), e["bas"], e["tur"]))
    return {"takvim": takvim, "yaklasan": yaklasan, "hasta_sayisi": len(hastalar), "uyarilar": uy}


def _ics_metin(s: str) -> str:
    return s.replace("\\", "\\\\").replace(";", "\\;").replace(",", "\\,").replace("\n", "\\n")


def ics_yaz(yol: Path, etkinlikler: list[dict]) -> None:
    satir = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//Gerisi Bende//Hasta Takvimi//TR", "CALSCALE:GREGORIAN"]
    for i, e in enumerate(etkinlikler, 1):
        satir.append("BEGIN:VEVENT")
        satir.append(f"UID:hasta-takvim-{i}-{e['bas']:%Y%m%d%H%M}@workers-workless")
        satir.append(f"DTSTAMP:{e['bas']:%Y%m%dT%H%M%S}Z")
        if e["saatli"]:
            satir.append(f"DTSTART:{e['bas']:%Y%m%dT%H%M%S}")
            if e["bit"]:
                satir.append(f"DTEND:{e['bit']:%Y%m%dT%H%M%S}")
        else:
            satir.append(f"DTSTART;VALUE=DATE:{e['bas']:%Y%m%d}")
            satir.append(f"DTEND;VALUE=DATE:{((e['bit'] or e['bas']).date() + timedelta(days=0 if e['bit'] and e['bit'].date() > e['bas'].date() else 1)):%Y%m%d}")
        satir.append(f"SUMMARY:{_ics_metin(e['hasta'] + ' – ' + e['tur'])}")
        if e["yer"]:
            satir.append(f"LOCATION:{_ics_metin(e['yer'])}")
        if e["aciklama"]:
            satir.append(f"DESCRIPTION:{_ics_metin(e['aciklama'])}")
        satir.append("END:VEVENT")
    satir.append("END:VCALENDAR")
    yol.parent.mkdir(parents=True, exist_ok=True)
    yol.write_text("\r\n".join(satir) + "\r\n", encoding="utf-8")


TUR_RENK = {"Operasyon": KIRMIZI, "Muayene": SARI, "Tetkik": SARI, "Kontrol": SARI, "Yatış": KIRMIZI, "Varış uçuşu": MAVI, "Dönüş uçuşu": MAVI}


def _zaman_yaz(dt: datetime | None, saatli: bool) -> str:
    return "" if dt is None else f"{dt:%d.%m.%Y %H:%M}" if saatli else f"{dt:%d.%m.%Y}"


def rapor_yaz(cikti: Path, r: dict, gun: date) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Takvim"
    bas = ["Hasta", "Etkinlik", "Başlangıç", "Bitiş", "Yer", "Açıklama"]
    _baslik(ws, bas, (20, 14, 17, 17, 30, 36), "C2")
    for e in r["takvim"]:
        _satir(ws, [e["hasta"], e["tur"], _zaman_yaz(e["bas"], e["saatli"]), _zaman_yaz(e["bit"], e["bit_saatli"]), e["yer"], e["aciklama"]], dolgu={2: TUR_RENK.get(e["tur"])}, sar=(6,))
    ws = wb.create_sheet("Yaklaşanlar")
    _baslik(ws, bas, (20, 14, 17, 17, 30, 36))
    for e in r["yaklasan"]:
        _satir(ws, [e["hasta"], e["tur"], _zaman_yaz(e["bas"], e["saatli"]), _zaman_yaz(e["bit"], e["bit_saatli"]), e["yer"], e["aciklama"]], dolgu={2: TUR_RENK.get(e["tur"])}, sar=(6,))
    ws.append([])
    ws.append([f"Rapor tarihi {gun:%d.%m.%Y}. Hatırlatma mesajlarını hastanın diline ve tercih ettiği kanala göre siz gönderin."])
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, gun: date, ileri: int = 3, ucus_sure: int = 0) -> dict:
    etkinlikler, uy = oku(yol)
    r = kontrol(etkinlikler, gun, ileri, ucus_sure)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, gun)
    ics_yaz(cikti.with_suffix(".ics"), r["takvim"])
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Uçuş, konaklama, muayene ve operasyon tarihlerinden hasta bazında takvim ve hatırlatmalar oluşturur.")
    p.add_argument("--etkinlikler", type=Path, help="Hasta, Tür, Başlangıç (tarih saat), Bitiş, Yer, Açıklama")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 09.10.2026)")
    p.add_argument("--gun", type=int, default=3, help="Hatırlatma listesine girecek gün sayısı (varsayılan 3)")
    p.add_argument("--ucus-sure", type=int, default=0, help="Operasyon ile dönüş uçuşu arasında olması istenen en az gün (varsayılan 0: kontrol yok)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "hasta_takvimi.xlsx")
    a = p.parse_args(argv)
    ornek = a.etkinlikler is None
    yol = a.etkinlikler or ORNEK / "hasta_etkinlikleri.csv"
    gun = tarih(a.tarih) if a.tarih else (date(2026, 10, 9) if ornek else date.today())
    if ornek and not a.ucus_sure:
        a.ucus_sure = 5
    if gun is None or a.gun < 0 or a.ucus_sure < 0:
        print("[X] --tarih GG.AA.YYYY, --gun ve --ucus-sure sıfır ya da üstü olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, gun, a.gun, a.ucus_sure)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {r['hasta_sayisi']} hasta · {len(r['takvim'])} etkinlik · {a.gun} gün içinde {len(r['yaklasan'])} etkinlik")
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[OK] Takvim: {a.cikti.with_suffix('.ics').resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
