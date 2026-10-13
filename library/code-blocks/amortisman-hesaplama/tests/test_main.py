import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def d(maliyet, omur, yontem="Normal", tarih=date(2024, 3, 15), kist=False, ad="X"):
    return {"ad": ad, "kod": "", "tarih": tarih, "maliyet": D(maliyet), "omur": omur, "yontem": yontem, "kist": kist}


class PlanTesti(unittest.TestCase):
    def test_normal(self):
        self.assertEqual(main.plan(d(120000, 5)), [D("24000.00")] * 5)
        self.assertEqual(main.plan(d(100000, 3)), [D("33333.33"), D("33333.33"), D("33333.34")])
        self.assertEqual((main.plan(d(1000, 1)), sum(main.plan(d("999.99", 7)))), ([D(1000)], D("999.99")))

    def test_azalan_bakiyeler(self):
        self.assertEqual(main.plan(d(100000, 5, "Azalan")), [D("40000.00"), D("24000.00"), D("14400.00"), D("8640.00"), D("12960.00")])
        self.assertEqual(main.plan(d(100000, 3, "Azalan")), [D("50000.00"), D("25000.00"), D("25000.00")])                    # 2 ÷ 3 = %66,67 → %50 sınırı
        self.assertEqual((main.plan(d(100000, 2, "Azalan")), main.plan(d(100000, 1, "Azalan")), main.plan(d(100000, 10, "Azalan"))[:3], sum(main.plan(d(100000, 10, "Azalan")))),
                         ([D("50000.00"), D("50000.00")], [D(100000)], [D("20000.00"), D("16000.00"), D("12800.00")], D(100000)))

    def test_kist(self):
        self.assertEqual(main.plan(d(600000, 5, tarih=date(2025, 9, 10), kist=True)), [D("40000.00"), D("120000.00"), D("120000.00"), D("120000.00"), D("200000.00")])
        self.assertEqual((main.plan(d(600000, 5, tarih=date(2025, 1, 31), kist=True)), main.plan(d(120000, 2, tarih=date(2025, 12, 1), kist=True)), main.plan(d(1000, 1, tarih=date(2025, 7, 1), kist=True))),
                         ([D("120000.00")] * 5, [D("5000.00"), D("115000.00")], [D(1000)]))
        self.assertEqual(main.plan(d(100000, 5, "Azalan", date(2025, 7, 1), True)), [D("20000.00"), D("24000.00"), D("14400.00"), D("8640.00"), D("32960.00")])


class HesapTesti(unittest.TestCase):
    def test_yil_secimi(self):
        h = lambda yil, **kw: (lambda x: (x["donem"], x["birikmis"], x["net"], x["durum"]))(main.hesapla([d(120000, 5, **kw)], yil)["demirbaslar"][0])   # noqa: E731
        self.assertEqual([h(y) for y in (2023, 2024, 2026, 2028, 2029)],
                         [(D(0), D(0), D(120000), "Aktife girmemiş"), (D("24000.00"), D("24000.00"), D("96000.00"), "Devam ediyor"), (D("24000.00"), D("72000.00"), D("48000.00"), "Devam ediyor"),
                          (D("24000.00"), D("120000.00"), D("0.00"), "Son yıl"), (D(0), D(120000), D(0), "Tamamen amorti edilmiş")])
        x = main.hesapla([d(100000, 5, "Azalan")], 2026)["demirbaslar"][0]
        self.assertEqual((x["oran"], x["donem"], x["birikmis"], x["net"], x["bitis"]), (D("40.00"), D("14400.00"), D("78400.00"), D("21600.00"), 2028))
        self.assertEqual((main.hesapla([d(100000, 3, "Azalan")], 2024)["demirbaslar"][0]["oran"], main.hesapla([d(100000, 8)], 2024)["demirbaslar"][0]["oran"]), (D("50.00"), D("12.50")))

    def test_toplam_ve_uyarilar(self):
        r = main.hesapla([d(120000, 5, ad="A"), d(100000, 5, "Azalan", date(2024, 1, 10), ad="B"), d(600000, 5, tarih=date(2025, 9, 10), kist=True, ad="C"), d(50000, 5, tarih=date(2020, 6, 1), ad="D"), d(30000, 3, tarih=date(2027, 2, 5), ad="E"),
                          d(1000, 1, tarih=date(2026, 5, 1), kist=True, ad="F"), d(1000, 4, "Azalan", date(2026, 7, 1), True, "G")], 2026)
        self.assertEqual((r["toplam"], [(o["yontem"], o["adet"], o["donem"]) for o in r["yontemler"]]),
                         ({"maliyet": D(902000), "donem": D("159650.00"), "birikmis": D("361650.00"), "net": D("540350.00")}, [("Azalan", 2, D("14650.00")), ("Normal", 5, D("145000.00"))]))
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]), [("Hesap yılından sonra aktife giren demirbaş", "E"), ("Kıst uygulanamadı", "F"), ("Kıst ve azalan bakiyeler birlikte", "G")])
        self.assertEqual([(p["yil"], p["tutar"], p["birikmis"], p["net"]) for p in r["planlar"] if p["ad"] == "C"],
                         [(2025, D("40000.00"), D("40000.00"), D("560000.00")), (2026, D("120000.00"), D("160000.00"), D("440000.00")), (2027, D("120000.00"), D("280000.00"), D("320000.00")),
                          (2028, D("120000.00"), D("400000.00"), D("200000.00")), (2029, D("200000.00"), D("600000.00"), D("0.00"))])
        with self.assertRaises(ValueError):
            main.hesapla([], 2026)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "a.xlsx"
            r = main.calistir(ORNEK / "demirbaslar.csv", cikti, 2026)
            self.assertEqual([(x["kod"], x["donem"], x["birikmis"], x["net"], x["durum"]) for x in r["demirbaslar"]],
                             [("DB-001", D("24000.00"), D("72000.00"), D("48000.00"), "Devam ediyor"), ("DB-002", D("14400.00"), D("78400.00"), D("21600.00"), "Devam ediyor"),
                              ("DB-003", D("120000.00"), D("160000.00"), D("440000.00"), "Devam ediyor"), ("DB-004", D(0), D("50000.00"), D("0.00"), "Tamamen amorti edilmiş"), ("DB-005", D(0), D(0), D("30000.00"), "Aktife girmemiş")])
            self.assertEqual((r["toplam"], [u["tur"] for u in r["uyarilar"]]),
                             ({"maliyet": D("900000.00"), "donem": D("158400.00"), "birikmis": D("360400.00"), "net": D("539600.00")}, ["Hesap yılından sonra aktife giren demirbaş"]))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Dönem Amortismanı", "Amortisman Planı", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "d.csv"
            y.write_text("Sabit Kıymet;Alış Tarihi;Alış Bedeli;Ömür;Amortisman Yöntemi;Binek\nA;15.03.2024;1.250.000,50;5;;\nB;15.03.2024;100;4;Azalan Bakiyeler;x\nC;x;100;5;;\nD;15.03.2024;0;5;;\nE;15.03.2024;100;2,5;;\n"
                         "F;15.03.2024;100;5;Üretim;\nG;15.03.2024;100;0;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["ad"], x["maliyet"], x["omur"], x["yontem"], x["kist"]) for x in l], [u["tur"] for u in uy]),
                             ([("A", D("1250000.50"), 5, "Normal", False), ("B", D(100), 4, "Azalan", True)], ["Demirbaş okunamadı"] * 5))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--demirbaslar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--yil", "2026"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(arg + ["--yil", "1800"]), 2)
            self.assertEqual(main.main(["--demirbaslar", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Demirbaş;Maliyet\nA;1\n", "Demirbaş;Aktife Giriş Tarihi;Maliyet;Faydalı Ömür\nA;x;1;1\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
