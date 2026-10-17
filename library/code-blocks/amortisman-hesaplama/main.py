"""
Amortisman Hesaplama — Gerisi Bende kod bloğu
Muhasebe › Genel Muhasebe Uzmanı

Demirbaş listesinden faydalı ömür ve yönteme göre seçilen yılın amortismanını, birikmiş amortismanı ve net defter değerini hesaplar.
  Normal amortisman      = maliyet ÷ faydalı ömür (yıl); her yıl eşit
  Azalan bakiyeler       = oran = min(2 ÷ faydalı ömür, %50);  yıl amortismanı = net defter değeri × oran;  son yıl kalan değerin tamamı
  Aktife giriş yılı      = tam yıl amortisman ayrılır (kıst işaretli olanlar hariç)
  Kıst amortisman        = "Kıst" sütunu Evet ise ilk yıl = yıllık tutar × (aktife girdiği ay dahil yıl sonuna kalan ay) ÷ 12; ayrılmayan kısım son yıla eklenir
  Birikmiş amortisman    = aktife giriş yılından --yil yılına kadar ayrılanların toplamı;   net defter değeri = maliyet − birikmiş
  Yuvarlama              = yıllık tutarlar kuruşa yuvarlanır; son yıl, toplam maliyete tam eşitlenecek şekilde kapatılır
Faydalı ömrü ve yöntemi siz verirsiniz; kod faydalı ömür tebliğlerini bilmez. Yeniden değerleme, enflasyon düzeltmesi, doğrudan gider yazma sınırı, binek otomobil gider kısıtlaması kapsam dışıdır.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"ad": ("demirbas", "demirbas adi", "sabit kiymet", "varlik", "kiymet adi"), "kod": ("kod", "demirbas no", "sabit kiymet no", "sicil no"), "tarih": ("aktife giris tarihi", "aktife giris", "alis tarihi", "aktiflestirme tarihi"),
         "maliyet": ("maliyet", "maliyet bedeli", "alis bedeli", "tutar"), "omur": ("faydali omur", "faydali omur yil", "omur", "omur yil"), "yontem": ("yontem", "amortisman yontemi"), "kist": ("kist", "kist amortisman", "binek")}
YONTEM = {"normal": "Normal", "dogrusal": "Normal", "esit": "Normal", "": "Normal", "azalan": "Azalan", "azalan bakiyeler": "Azalan", "hizlandirilmis": "Azalan"}
EVET = {"evet", "e", "x", "var", "1", "true"}
Z = Decimal(0)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for n, r in enumerate(kayitlar(yol, SUTUN, ("ad", "tarih", "maliyet", "omur")), 1):
        ad = metin(r["ad"])
        if not ad:
            continue
        t, m, o, y = tarih(r["tarih"]), para(r["maliyet"]), para(r["omur"]), YONTEM.get(katla(r["yontem"]))
        if t is None or m is None or m <= 0 or o is None or o < 1 or o != o.to_integral_value() or o > 100 or y is None:
            uy.append(uyari("Yüksek", "Demirbaş okunamadı", f"{ad} (satır {r['_satir']})", f"Aktife giriş '{metin(r['tarih'])}', maliyet '{metin(r['maliyet'])}', faydalı ömür '{metin(r['omur'])}', yöntem '{metin(r['yontem'])}'. "
                            "Maliyet pozitif, ömür 1–100 arası tam yıl, yöntem Normal veya Azalan olmalı. Satır hesaplanmadı."))
            continue
        l.append({"ad": ad, "kod": metin(r["kod"]), "tarih": t, "maliyet": m, "omur": int(o), "yontem": y, "kist": katla(r["kist"]) in EVET})
    return l, uy


def plan(d: dict) -> list[Decimal]:
    """Aktife giriş yılından başlayarak yıllık amortisman tutarları (uzunluk = faydalı ömür)."""
    m, n = d["maliyet"], d["omur"]
    if d["yontem"] == "Normal":
        p = [m / n] * n
    else:
        oran, p, kalan = min(Decimal(2) / n, Decimal("0.5")), [], m
        for i in range(n):
            a = kalan if i == n - 1 else kalan * oran
            p.append(a)
            kalan -= a
    if d["kist"] and n > 1:
        ertelenen = p[0] * (d["tarih"].month - 1) / 12
        p[0] -= ertelenen
        p[-1] += ertelenen
    p = [yuvarla(a) for a in p]
    p[-1] = m - sum(p[:-1], Z)
    return p


def hesapla(demirbaslar: list[dict], yil: int) -> dict:
    if not demirbaslar:
        raise ValueError("Hesaplanacak demirbaş yok.")
    uy, liste, planlar = [], [], []
    for d in demirbaslar:
        p, bas = plan(d), d["tarih"].year
        idx = yil - bas
        if idx < 0:
            donem, birikmis, durum = Z, Z, "Aktife girmemiş"
            uy.append(uyari("Bilgi", "Hesap yılından sonra aktife giren demirbaş", d["ad"], f"Aktife giriş {d['tarih']:%d.%m.%Y}; {yil} yılı için amortisman hesaplanmadı."))
        elif idx >= d["omur"]:
            donem, birikmis, durum = Z, d["maliyet"], "Tamamen amorti edilmiş"
        else:
            donem, birikmis, durum = p[idx], sum(p[:idx + 1], Z), "Son yıl" if idx == d["omur"] - 1 else "Devam ediyor"
        if d["kist"] and d["yontem"] == "Azalan":
            uy.append(uyari("Orta", "Kıst ve azalan bakiyeler birlikte", d["ad"], "Kıst amortismanda ilk yıl ayrılmayan kısım son yıla eklendi; ara yıllar tam yıl gibi hesaplandı. Bu birleşimin uygulamasını meslek mensubunuzla doğrulayın."))
        if d["kist"] and d["omur"] == 1:
            uy.append(uyari("Orta", "Kıst uygulanamadı", d["ad"], "Faydalı ömür 1 yıl olduğu için kıst kısım ertelenecek yıl yok; maliyetin tamamı ilk yıla yazıldı."))
        oran = Decimal(100) / d["omur"] if d["yontem"] == "Normal" else min(Decimal(200) / d["omur"], Decimal(50))
        liste.append(d | {"oran": yuvarla(oran), "donem": donem, "birikmis": birikmis, "net": d["maliyet"] - birikmis, "durum": durum, "bitis": bas + d["omur"] - 1})
        kum = Z
        for i, a in enumerate(p):
            kum += a
            planlar.append({"ad": d["ad"], "kod": d["kod"], "yil": bas + i, "tutar": a, "birikmis": kum, "net": d["maliyet"] - kum})
    toplam = {a: sum((x[a] for x in liste), Z) for a in ("maliyet", "donem", "birikmis", "net")}
    yontem = {}
    for x in liste:
        o = yontem.setdefault(x["yontem"], {"yontem": x["yontem"], "adet": 0, "maliyet": Z, "donem": Z, "birikmis": Z, "net": Z})
        o["adet"] += 1
        for a in ("maliyet", "donem", "birikmis", "net"):
            o[a] += x[a]
    return {"demirbaslar": liste, "planlar": planlar, "toplam": toplam, "yontemler": sorted(yontem.values(), key=lambda o: o["yontem"]), "yil": yil, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb, t = Workbook(), r["toplam"]
    ws = wb.active
    ws.title = "Dönem Amortismanı"
    _baslik(ws, ["Kod", "Demirbaş", "Aktife Giriş", "Maliyet", "Faydalı Ömür (yıl)", "Yöntem", "Oran %", "Kıst", f"{r['yil']} Amortismanı", "Birikmiş Amortisman", "Net Defter Değeri", "Son Amortisman Yılı", "Durum"],
            (12, 34, 13, 16, 9, 10, 9, 7, 16, 17, 16, 10, 22), "C2")
    renk = {"Tamamen amorti edilmiş": MAVI, "Son yıl": SARI, "Aktife girmemiş": SARI}
    for x in r["demirbaslar"]:
        _satir(ws, [x["kod"], x["ad"], x["tarih"], x["maliyet"], x["omur"], x["yontem"], x["oran"], "Evet" if x["kist"] else "", x["donem"], x["birikmis"], x["net"], x["bitis"], x["durum"]], para_s=(4, 7, 9, 10, 11), tarih_s=(3,),
               dolgu={13: renk.get(x["durum"])})
    n = _satir(ws, ["", "Toplam", None, t["maliyet"], None, "", None, "", t["donem"], t["birikmis"], t["net"]], para_s=(4, 9, 10, 11))
    for c in ws[n]:
        c.font = Font(bold=True)
    ws.append([])
    for o in r["yontemler"]:
        _satir(ws, ["", f"{o['yontem']} ({o['adet']} adet)", None, o["maliyet"], None, "", None, "", o["donem"], o["birikmis"], o["net"]], para_s=(4, 9, 10, 11))
    ws.append([])
    ws.append(["Normal: maliyet ÷ ömür. Azalan bakiyeler: oran normalin iki katı, en çok %50; son yıl kalan tamamı. Aktife giriş yılında tam yıl amortisman ayrılır; 'Kıst' işaretlilerde ilk yıl ay kesriyle, kalan son yıla eklenir."])
    ws = wb.create_sheet("Amortisman Planı")
    _baslik(ws, ["Kod", "Demirbaş", "Yıl", "Amortisman", "Birikmiş Amortisman", "Net Defter Değeri"], (12, 34, 8, 16, 18, 17))
    for p in r["planlar"]:
        _satir(ws, [p["kod"], p["ad"], p["yil"], p["tutar"], p["birikmis"], p["net"]], para_s=(4, 5, 6), dolgu={3: SARI if p["yil"] == r["yil"] else None})
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, yil: int) -> dict:
    demirbaslar, uy = oku(yol)
    r = hesapla(demirbaslar, yil)
    uy += r["uyarilar"]
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Demirbaş listesinden yıl amortismanını, birikmiş amortismanı ve net defter değerini hesaplar.")
    p.add_argument("--demirbaslar", type=Path, default=ORNEK / "demirbaslar.csv", help="Demirbaş, Kod, Aktife Giriş Tarihi, Maliyet, Faydalı Ömür (yıl), Yöntem (Normal / Azalan), Kıst")
    p.add_argument("--yil", type=int, help="Amortismanı hesaplanacak yıl (varsayılan: içinde bulunulan yıl)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "amortisman.xlsx")
    a = p.parse_args(argv)
    yil = a.yil or (2026 if a.demirbaslar == ORNEK / "demirbaslar.csv" else date.today().year)
    if not 1950 <= yil <= 2200:
        print("[X] --yil geçerli bir yıl olmalı")
        return 2
    if not a.demirbaslar.exists():
        print(f"[X] Dosya bulunamadı: {a.demirbaslar}")
        return 1
    try:
        r = calistir(a.demirbaslar, a.cikti, yil)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {len(r['demirbaslar'])} demirbaş · maliyet {t['maliyet']:,.2f} · {yil} amortismanı {t['donem']:,.2f} · birikmiş {t['birikmis']:,.2f} · net defter değeri {t['net']:,.2f}")
    for u in r["uyarilar"]:
        if u["onem"] != "Bilgi":
            print(f"[!] {u['tur']}: {u['kim']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
