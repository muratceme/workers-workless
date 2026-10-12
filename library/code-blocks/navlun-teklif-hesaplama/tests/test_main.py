import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def kp(adet, boy, en, yukseklik, kg):
    return {"ad": "K", "adet": adet, "boy": D(boy), "en": D(en), "yukseklik": D(yukseklik), "kg": D(kg)}


def tr(birim, fiyat, asgari=None, para="USD", kalem="X"):
    return {"kalem": kalem, "birim": birim, "fiyat": D(fiyat), "asgari": None if asgari is None else D(asgari), "para": para}


KAPLAR = lambda: [kp(2, 60, 50, 40, 25), kp(1, 120, 80, 110, 150)]   # noqa: E731


class HesapTesti(unittest.TestCase):
    def test_hava(self):
        r = main.hesapla(KAPLAR(), [tr("esas", "2.5", 100), tr("esas", "0.8"), tr("esas", "0.15", 25), tr("sabit", 35), tr("kap", 4)])
        self.assertEqual((r["brut"], r["hacimsel"], r["esas"], r["belirleyen"], r["m3"], r["adet"]), (D(200), D("216.00"), D("216.0"), "Hacim", D("1.296"), 3))
        self.assertEqual(([s["tutar"] for s in r["satirlar"]], r["toplam"]), ([D("540.00"), D("172.80"), D("32.40"), D("35.00"), D("12.00")], {"USD": D("792.20")}))
        self.assertEqual([main.hesapla([kp(1, 100, 100, c, 10)], [tr("sabit", 1)])["esas"] for c in (60, 61, 63, 64)], [D("100.0"), D("102.0"), D("105.0"), D("107.0")])      # 101,67 → 102,0; 106,67 → 107,0
        r = main.hesapla([kp(1, 10, 10, 10, "12.2")], [tr("esas", 3, 50)])
        self.assertEqual((r["esas"], r["belirleyen"], r["satirlar"][0]["ham"], r["satirlar"][0]["tutar"], r["satirlar"][0]["asgari_uygulandi"]), (D("12.5"), "Ağırlık", D("37.50"), D(50), True))
        self.assertEqual((main.hesapla(KAPLAR(), [tr("sabit", 1)], bolen=D(5000))["esas"], main.hesapla([kp(1, 100, 100, 61, 10)], [tr("sabit", 1)], yuvarlama=D(0))["esas"]), (D("259.5"), D("101.67")))

    def test_deniz_ve_toplamlar(self):
        r = main.hesapla(KAPLAR(), [tr("esas", 60, 60), tr("esas", 25, 40), tr("sabit", 50), tr("sabit", "1500", para="TRY")], "deniz")
        self.assertEqual((r["ton"], r["m3"], r["esas"], r["belirleyen"]), (D("0.200"), D("1.296"), D("1.296"), "Hacim"))
        self.assertEqual(([s["tutar"] for s in r["satirlar"]], r["toplam"]), ([D("77.76"), D("40"), D("50.00"), D("1500.00")], {"USD": D("167.76"), "TRY": D("1500.00")}))
        r = main.hesapla([kp(1, 100, 100, 100, 2500)], [tr("esas", 60)], "deniz")
        self.assertEqual((r["esas"], r["belirleyen"], r["toplam"]), (D("2.500"), "Ağırlık", {"USD": D("150.00")}))
        self.assertEqual([main.birim_coz(x) for x in ("Ücrete esas", "W/M", "kg başına", "Kap", "Sabit", "", "B/L", "günlük")], ["esas", "esas", "esas", "kap", "sabit", "sabit", "sabit", None])
        for k, t in (([], [tr("sabit", 1)]), (KAPLAR(), [])):
            with self.assertRaises(ValueError):
                main.hesapla(k, t)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "n.xlsx"
            r = main.calistir(ORNEK / "kaplar.csv", ORNEK / "tarife_hava.csv", cikti)
            self.assertEqual((r["esas"], r["toplam"], r["uyarilar"]), (D("216.0"), {"USD": D("780.20")}, []))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Teklif", "Kap Listesi", "Uyarılar"])
            r = main.calistir(ORNEK / "kaplar.csv", ORNEK / "tarife_deniz.csv", cikti, "deniz")
            self.assertEqual((r["esas"], r["toplam"], [s["asgari_uygulandi"] for s in r["satirlar"]]), (D("1.296"), {"USD": D("167.76")}, [False, True, False]))

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            k, t = Path(tmp) / "k.csv", Path(tmp) / "t.csv"
            k.write_text("Kap;Adet;Boy;En;Yükseklik;Ağırlık\nA;2;60;50;40;25,5\nB;1;120;80;;150\nC;1;1,2;0,8;1;10\n", encoding="utf-8")
            t.write_text("Kalem;Birim;Birim Fiyat;Asgari;Para Birimi\nNavlun;kg;2,50;100;usd\nDepo;günlük;5;;USD\nEvrak;;35;;\nBozuk;Sabit;çok;;USD\n", encoding="utf-8")
            kl, tl, uy = main.oku(k, t)
            self.assertEqual(([(x["ad"], x["adet"], x["kg"]) for x in kl], [(x["kalem"], x["birim"], x["fiyat"], x["asgari"], x["para"]) for x in tl]),
                             ([("A", 2, D("25.5")), ("C", 1, D(10))], [("Navlun", "esas", D("2.50"), D(100), "USD"), ("Evrak", "sabit", D(35), None, "")]))
            self.assertEqual(sorted(u["tur"] for u in uy), ["Kap satırı okunamadı", "Tarife kalemi okunamadı", "Tarife kalemi okunamadı", "Ölçü birimi şüpheli"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--kaplar", str(k), "--tarife", str(t), "--cikti", c]
            self.assertEqual(sorted(u["tur"] for u in main.calistir(k, t, Path(c))["uyarilar"])[:2], ["Birden çok para birimi", "Kap satırı okunamadı"])
            self.assertEqual(main.main(arg + ["--mod", "deniz"]), 0)
            self.assertEqual(main.main(arg + ["--bolen", "5000", "--yuvarlama", "1"]), 0)
            self.assertEqual(main.main(["--kaplar", str(k), "--cikti", c]), 2)                      # kendi kap dosyanızla tarife zorunlu
            self.assertEqual(main.main(arg + ["--bolen", "0"]), 2)
            self.assertEqual(main.main(arg[:3] + [str(Path(tmp) / "yok.csv")]), 1)
            t.write_text("Kalem;Birim Fiyat\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)
            self.assertEqual(main.main(["--cikti", c, "--mod", "deniz"]), 0)


if __name__ == "__main__":
    unittest.main()
