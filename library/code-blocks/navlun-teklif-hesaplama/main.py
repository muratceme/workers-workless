"""
Navlun Teklif Hesaplama — Workers / Workless kod bloğu
Deniz ve Hava Yolu Operasyon › Denizyolu Operasyon Uzmanı

Ölçü ve ağırlıktan hava kargoda ücrete esas ağırlığı (1:6000), denizyolunda W/M'yi hesaplayıp navlun ve ek masraflarla teklif oluşturur.
  Hava   : hacimsel ağırlık (kg) = Σ adet × boy × en × yükseklik (cm) ÷ 6000      (--bolen ile değişir)
           ücrete esas ağırlık   = en büyük(brüt ağırlık, hacimsel ağırlık), 0,5 kg'a yukarı yuvarlanır (--yuvarlama)
  Deniz  : hacim (m³) = Σ adet × boy × en × yükseklik ÷ 1.000.000;   ağırlık (ton) = brüt kg ÷ 1000
           W/M = en büyük(ton, m³)                                                (1 ton = 1 m³ kabulü; parsiyel / LCL yükler için)
  Tarife : her kalem "Ücrete esas" (birim fiyat × ücrete esas ağırlık veya W/M), "Kap" (× kap adedi) veya "Sabit" olarak hesaplanır; "Asgari" yazılmışsa tutar onun altına inmez
  Toplam : para birimi bazında ayrı ayrı; kur çevirisi yapılmaz
Birim fiyatlar ve ek masraflar sizin tarife dosyanızdan gelir; kod hiçbir hattın veya havayolunun fiyatını bilmez. Sonuç tekliftir, taşıyıcının tartı ve ölçümü esastır.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


KAP = {"ad": ("kap", "kalem", "aciklama", "yuk"), "adet": ("adet", "kap adedi"), "boy": ("boy", "boy cm", "uzunluk"), "en": ("en", "en cm", "genislik"), "yuk": ("yukseklik", "yukseklik cm"),
       "kg": ("agirlik", "brut agirlik", "agirlik kg", "kg", "birim agirlik")}
TARIFE = {"kalem": ("kalem", "masraf", "masraf kalemi", "aciklama"), "birim": ("birim", "hesap sekli", "esas"), "fiyat": ("birim fiyat", "fiyat", "tutar"), "asgari": ("asgari", "minimum", "asgari tutar"),
          "para": ("para birimi", "doviz", "pb")}
Z = Decimal(0)


def birim_coz(x) -> str | None:
    k = katla(x)
    if k in ("", "sabit", "sefer", "konsimento", "sevkiyat", "dosya", "awb", "b l"):
        return "sabit"
    if k in ("kap", "adet", "kap basina", "koli"):
        return "kap"
    if any(a in k for a in ("ucrete esas", "w m", "wm", "kg", "ton", "m3")):
        return "esas"
    return None


def oku(kap_yol: Path, tarife_yol: Path) -> tuple[list[dict], list[dict], list[dict]]:
    uy, kaplar, tarife = [], [], []
    for r in kayitlar(kap_yol, KAP, ("adet", "boy", "en", "yuk", "kg")):
        if not any(metin(r[a]) for a in ("adet", "boy", "en", "yuk", "kg")):
            continue
        ad = metin(r["ad"]) or f"Satır {r['_satir']}"
        v = {a: para(r[a]) for a in ("adet", "boy", "en", "yuk", "kg")}
        if any(x is None or x <= 0 for x in v.values()) or v["adet"] != int(v["adet"]):
            uy.append(uyari("Yüksek", "Kap satırı okunamadı", ad, "Adet (tam sayı), boy, en, yükseklik (cm) ve ağırlık (kg) pozitif olmalı. Satır alınmadı; teklif eksik hesaplanır."))
            continue
        if max(v["boy"], v["en"], v["yuk"]) < 5:
            uy.append(uyari("Orta", "Ölçü birimi şüpheli", ad, "Ölçüler santimetre olmalı (metre yazılmış olabilir)."))
        kaplar.append({"ad": ad, "adet": int(v["adet"]), "boy": v["boy"], "en": v["en"], "yukseklik": v["yuk"], "kg": v["kg"]})
    for r in kayitlar(tarife_yol, TARIFE, ("kalem", "fiyat")):
        k = metin(r["kalem"])
        if not k:
            continue
        f, asgari, b = para(r["fiyat"]), para(r["asgari"]), birim_coz(r["birim"])
        if f is None or f < 0 or b is None:
            uy.append(uyari("Yüksek", "Tarife kalemi okunamadı", k, f"Birim '{metin(r['birim'])}', birim fiyat '{metin(r['fiyat'])}'. Birim: Ücrete esas, Kap veya Sabit olmalı. Kalem teklife alınmadı."))
            continue
        tarife.append({"kalem": k, "birim": b, "fiyat": f, "asgari": asgari if asgari and asgari > 0 else None, "para": metin(r["para"]).upper()})
    return kaplar, tarife, uy


def hesapla(kaplar: list[dict], tarife: list[dict], mod: str = "hava", bolen: Decimal = Decimal(6000), yuvarlama: Decimal = Decimal("0.5")) -> dict:
    if not kaplar:
        raise ValueError("Hesaplanacak kap yok.")
    if not tarife:
        raise ValueError("Tarife dosyasında kalem yok.")
    for k in kaplar:
        k["hacim_cm3"] = k["adet"] * k["boy"] * k["en"] * k["yukseklik"]
        k["hacim"] = yuvarla(k["hacim_cm3"] / 1000000, "0.001")
        k["brut"] = k["adet"] * k["kg"]
        k["hacimsel"] = yuvarla(k["hacim_cm3"] / bolen)
    cm3, brut, adet = sum((k["hacim_cm3"] for k in kaplar), Z), sum((k["brut"] for k in kaplar), Z), sum(k["adet"] for k in kaplar)
    m3 = yuvarla(cm3 / 1000000, "0.001")
    if mod == "hava":
        hacimsel = yuvarla(cm3 / bolen)
        ham = max(brut, hacimsel)
        esas = Decimal(math.ceil(ham / yuvarlama)) * yuvarlama if yuvarlama else ham
        olcu = {"hacimsel": hacimsel, "esas": esas, "birim": "kg", "belirleyen": "Hacim" if hacimsel > brut else "Ağırlık"}
    else:
        ton = yuvarla(brut / 1000, "0.001")
        olcu = {"ton": ton, "esas": max(ton, m3), "birim": "W/M", "belirleyen": "Hacim" if m3 > ton else "Ağırlık"}
    satirlar, toplam = [], defaultdict(lambda: Z)
    for t in tarife:
        carpan = {"esas": olcu["esas"], "kap": Decimal(adet), "sabit": Decimal(1)}[t["birim"]]
        ham = yuvarla(t["fiyat"] * carpan)
        tutar = max(ham, t["asgari"]) if t["asgari"] else ham
        satirlar.append(t | {"carpan": carpan, "ham": ham, "tutar": tutar, "asgari_uygulandi": tutar > ham})
        toplam[t["para"]] += tutar
    return {"kaplar": kaplar, "satirlar": satirlar, "toplam": dict(toplam), "mod": mod, "adet": adet, "brut": brut, "m3": m3, "ayar": (bolen, yuvarlama)} | olcu


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    hava = r["mod"] == "hava"
    ws = wb.active
    ws.title = "Teklif"
    _baslik(ws, ["Masraf Kalemi", "Hesap Şekli", "Birim Fiyat", "Çarpan", "Asgari", "Tutar", "Para Birimi", "Not"], (36, 16, 12, 12, 12, 14, 10, 22))
    ad = {"esas": f"Ücrete esas ({r['birim']})", "kap": "Kap başına", "sabit": "Sabit"}
    for s in r["satirlar"]:
        _satir(ws, [s["kalem"], ad[s["birim"]], s["fiyat"], s["carpan"], s["asgari"], s["tutar"], s["para"], "Asgari tutar uygulandı" if s["asgari_uygulandi"] else ""], para_s=(3, 5, 6),
               dolgu={8: SARI if s["asgari_uygulandi"] else None})
    for pb, t in r["toplam"].items():
        n = _satir(ws, [f"TOPLAM {pb}".strip(), "", None, None, None, t, pb], para_s=(6,), dolgu={6: YESIL})
        for c in ws[n]:
            c.font = Font(bold=True)
    ws.append([])
    b, y = r["ayar"]
    ws.append([(f"Hava kargo: brüt {r['brut']:f} kg, hacimsel {r['hacimsel']:f} kg (÷ {b:f}), ücrete esas {r['esas']:f} kg ({y:f} kg'a yukarı yuvarlandı)." if hava else
                f"Denizyolu: {r['ton']:f} ton, {r['m3']:f} m³, W/M {r['esas']:f}.") + f" Belirleyen: {r['belirleyen']}. Taşıyıcının tartı ve ölçümü esastır; fiyatlar tarife dosyanızdan alınmıştır."])
    ws = wb.create_sheet("Kap Listesi")
    _baslik(ws, ["Kap", "Adet", "Boy (cm)", "En (cm)", "Yükseklik (cm)", "Birim Ağırlık (kg)", "Brüt Ağırlık (kg)", "Hacim (m³)", *(["Hacimsel Ağırlık (kg)"] if hava else [])], (30, 8, 9, 9, 10, 11, 12, 11, 12))
    for k in r["kaplar"]:
        _satir(ws, [k["ad"], k["adet"], k["boy"], k["en"], k["yukseklik"], k["kg"], k["brut"], k["hacim"], *([k["hacimsel"]] if hava else [])], para_s=(7, 9))
        ws.cell(ws.max_row, 8).number_format = "0.000"
    _satir(ws, ["Toplam", r["adet"], "", "", "", "", r["brut"], r["m3"], *([r["hacimsel"]] if hava else [])], para_s=(7, 9))
    ws.cell(ws.max_row, 8).number_format = "0.000"
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(kap_yol: Path, tarife_yol: Path, cikti: Path, mod: str = "hava", **ayar) -> dict:
    kaplar, tarife, uy = oku(kap_yol, tarife_yol)
    r = hesapla(kaplar, tarife, mod, **ayar)
    if len(r["toplam"]) > 1:
        uy.append(uyari("Bilgi", "Birden çok para birimi", tarife_yol.name, f"Toplamlar ayrı verildi: {', '.join(p or '(boş)' for p in r['toplam'])}. Kur çevirisi yapılmaz."))
    if "" in r["toplam"]:
        uy.append(uyari("Orta", "Para birimi yok", tarife_yol.name, "Bazı tarife kalemlerinde para birimi boş."))
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Hava kargoda ücrete esas ağırlık, denizyolunda W/M hesaplayıp navlun teklifi oluşturur.")
    p.add_argument("--kaplar", type=Path, default=ORNEK / "kaplar.csv", help="Kap, Adet, Boy (cm), En (cm), Yükseklik (cm), Ağırlık (kg, kap başına)")
    p.add_argument("--tarife", type=Path, help="Kalem, Birim (Ücrete esas / Kap / Sabit), Birim Fiyat, Asgari, Para Birimi")
    p.add_argument("--mod", choices=("hava", "deniz"), default="hava", help="Taşıma türü (varsayılan hava)")
    p.add_argument("--bolen", default="6000", help="Hava kargoda hacimsel ağırlık böleni (varsayılan 6000)")
    p.add_argument("--yuvarlama", default="0,5", help="Hava kargoda ücrete esas ağırlığın yukarı yuvarlanacağı adım, kg (varsayılan 0,5; 0 = yuvarlama yok)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "navlun_teklifi.xlsx")
    a = p.parse_args(argv)
    tarife = a.tarife or (ORNEK / f"tarife_{a.mod}.csv" if a.kaplar == ORNEK / "kaplar.csv" else None)
    b, y = para(a.bolen), para(a.yuvarlama)
    if tarife is None or b is None or b <= 0 or y is None or y < 0:
        print("[X] --tarife verilmeli; --bolen pozitif, --yuvarlama sıfır veya pozitif sayı olmalı")
        return 2
    for d in (a.kaplar, tarife):
        if not d.exists():
            print(f"[X] Dosya bulunamadı: {d}")
            return 1
    try:
        r = calistir(a.kaplar, tarife, a.cikti, a.mod, bolen=b, yuvarlama=y)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    if r["mod"] == "hava":
        print(f"[OK] Hava: {r['adet']} kap · brüt {r['brut']:f} kg · hacimsel {r['hacimsel']:f} kg · ücrete esas {r['esas']:f} kg ({r['belirleyen'].lower()} belirledi)")
    else:
        print(f"[OK] Deniz: {r['adet']} kap · {r['ton']:f} ton · {r['m3']:f} m³ · W/M {r['esas']:f} ({r['belirleyen'].lower()} belirledi)")
    for pb, t in r["toplam"].items():
        print(f"[OK] Teklif toplamı: {t:,.2f} {pb}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
