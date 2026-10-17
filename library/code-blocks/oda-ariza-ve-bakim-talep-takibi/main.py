"""
Oda Arıza ve Bakım Talep Takibi — Gerisi Bende kod bloğu
Kat Hizmetleri › Kat Hizmetleri Müdürü (Housekeeper)
Odalardan bildirilen arıza ve bakım taleplerini oda, kategori, öncelik ve bekleme süresine göre takip eder.
  Süre (saat)       = tamamlanma − bildirim; açık talepte rapor zamanı − bildirim
  Hedef süre        = önceliğe göre: Acil --acil (4) · Normal --normal (24) · Düşük --dusuk (72) saat
  Hedefte %         = hedef süresi içinde tamamlanan ÷ tamamlanan × 100
  Satışa kapalı oda = "Satışa Kapalı" işaretli talepler; kayıp oda-gece = kapalı kalınan takvim günü (en az 1)
  Tekrar eden arıza = aynı oda + kategoride --gun (30) gün içinde --tekrar (3) ya da daha fazla talep
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
SUTUN = {"no": ("talep no", "is emri no", "kayit no", "no"), "oda": ("oda", "oda no", "lokasyon", "yer"), "bildirim": ("bildirim zamani", "bildirim tarihi", "bildirim", "acilis", "tarih"),
         "kategori": ("kategori", "ariza turu", "tur", "konu"), "aciklama": ("aciklama", "ariza", "detay"), "oncelik": ("oncelik", "aciliyet"),
         "durum": ("durum", "statu"), "bitis": ("tamamlanma zamani", "tamamlanma tarihi", "tamamlanma", "kapanis", "bitis"), "kapali": ("satisa kapali", "ooo", "out of order", "kapali"),
         "sorumlu": ("sorumlu", "teknisyen", "atanan")}
ONCELIK = ("Acil", "Normal", "Düşük")


def zaman(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    if isinstance(x, date):
        return datetime(x.year, x.month, x.day)
    s = " ".join(str(x or "").strip().split())
    for f in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d/%m/%Y %H:%M", "%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def _oncelik(x) -> str | None:
    k = katla(x)
    if not k or k in ("normal", "orta", "standart"):
        return "Normal"
    if k in ("acil", "yuksek", "kritik", "cok acil", "urgent", "1"):
        return "Acil"
    if k in ("dusuk", "planli", "ertelenebilir", "low", "3"):
        return "Düşük"
    return None


def _durum(x) -> str | None:
    k = katla(x)
    if not k or k in ("acik", "yeni", "bekliyor", "beklemede", "atandi"):
        return "Açık"
    if k in ("devam", "devam ediyor", "islemde", "parca bekliyor", "calisiliyor"):
        return "Devam ediyor"
    if k in ("tamamlandi", "kapandi", "kapali", "bitti", "cozuldu", "yapildi"):
        return "Tamamlandı"
    if "iptal" in k:
        return "İptal"
    return None


def _saat(a: datetime, b: datetime) -> Decimal:
    return yuvarla(Decimal((b - a).total_seconds()) / 3600, "0.1")


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("oda", "bildirim")):
        if not metin(r["oda"]) and not metin(r["bildirim"]):
            continue
        no = metin(r["no"])
        kim = no or f"Oda {metin(r['oda'])}"
        b, d, o = zaman(r["bildirim"]), _durum(r["durum"]), _oncelik(r["oncelik"])
        bit = zaman(r["bitis"]) if metin(r["bitis"]) else None
        if b is None or d is None or o is None or not metin(r["oda"]) or (metin(r["bitis"]) and bit is None):
            uy.append(uyari("Yüksek", "Satır okunamadı", kim, f"Satır {r['_satir']}: oda, bildirim zamanı, öncelik ('{metin(r['oncelik'])}'), durum ('{metin(r['durum'])}') ya da tamamlanma okunamadı."))
            continue
        if no and katla(no) in gorulen:
            uy.append(uyari("Orta", "Talep iki kez yazılmış", kim, f"Satır {r['_satir']}: ikinci satır alınmadı."))
            continue
        if no:
            gorulen.add(katla(no))
        if d == "İptal":
            continue
        if bit and bit < b:
            uy.append(uyari("Yüksek", "Tamamlanma bildirimden önce", kim, f"Bildirim {b:%d.%m.%Y %H:%M}, tamamlanma {bit:%d.%m.%Y %H:%M}; tamamlanma yok sayıldı."))
            bit = None
        if bit and d != "Tamamlandı":
            d = "Tamamlandı"
        if d == "Tamamlandı" and bit is None:
            uy.append(uyari("Bilgi", "Tamamlanma zamanı yok", kim, "Talep tamamlandı görünüyor ama tamamlanma zamanı yazılmamış; süre hesabına girmedi."))
        l.append({"no": no, "oda": metin(r["oda"]), "bildirim": b, "kategori": metin(r["kategori"]) or "Belirtilmemiş", "aciklama": metin(r["aciklama"]), "oncelik": o, "durum": d,
                  "bitis": bit, "kapali": katla(r["kapali"]) in ("evet", "e", "x", "var", "ooo", "1", "kapali"), "sorumlu": metin(r["sorumlu"])})
    return l, uy


def analiz(talepler: list[dict], simdi: datetime, hedef: dict, tekrar: int = 3, gun: int = 30) -> dict:
    if not talepler:
        raise ValueError("Okunabilen talep yok.")
    uy = []
    for t in talepler:
        if t["bildirim"] > simdi:
            t["sure"] = None
        else:
            son = t["bitis"] or (simdi if t["durum"] != "Tamamlandı" else None)
            t["sure"] = _saat(t["bildirim"], son) if son else None
        t["hedef"] = hedef[t["oncelik"]]
        t["asim"] = t["sure"] is not None and t["sure"] > t["hedef"]
        t["acik"] = t["durum"] != "Tamamlandı"
        if t["kapali"]:
            son = t["bitis"] or simdi
            t["kayip_gece"] = max((son.date() - t["bildirim"].date()).days, 1) if son >= t["bildirim"] else 0
        else:
            t["kayip_gece"] = 0
        kim = f"{t['no']} · Oda {t['oda']}" if t["no"] else f"Oda {t['oda']}"
        if t["acik"] and t["asim"]:
            uy.append(uyari("Yüksek" if t["oncelik"] == "Acil" else "Orta", "Hedef süre aşıldı", kim,
                            f"{t['kategori']} ({t['oncelik']}): {t['sure']} saattir açık, hedef {t['hedef']} saat. {t['aciklama']}".strip()))
        if t["acik"] and t["kapali"]:
            uy.append(uyari("Orta", "Oda satışa kapalı", kim, f"{t['kategori']}: {t['bildirim']:%d.%m.%Y %H:%M} tarihinden beri satışa kapalı ({t['kayip_gece']} gece)."))
    gruplar = {}
    for t in talepler:
        gruplar.setdefault((katla(t["oda"]), katla(t["kategori"])), []).append(t)
    tekrarlar = []
    for l in gruplar.values():
        l.sort(key=lambda t: t["bildirim"])
        en_cok = max(sum(1 for y in l if x["bildirim"] <= y["bildirim"] <= x["bildirim"] + timedelta(days=gun)) for x in l)
        if en_cok >= tekrar:
            tekrarlar.append({"oda": l[0]["oda"], "kategori": l[0]["kategori"], "adet": en_cok, "toplam": len(l), "ilk": l[0]["bildirim"], "son": l[-1]["bildirim"]})
            uy.append(uyari("Orta", "Tekrar eden arıza", f"Oda {l[0]['oda']}", f"{l[0]['kategori']}: {gun} gün içinde {en_cok} talep (toplam {len(l)}). Kalıcı onarım / yenileme değerlendirin."))

    def ozet(fn, sira=None):
        g = {}
        for t in talepler:
            g.setdefault(fn(t), []).append(t)
        satir = []
        for ad, l in g.items():
            biten = [t for t in l if not t["acik"] and t["sure"] is not None]
            satir.append({"ad": ad, "adet": len(l), "acik": sum(1 for t in l if t["acik"]), "geciken": sum(1 for t in l if t["acik"] and t["asim"]), "tamam": len(biten),
                          "ort": yuvarla(sum((t["sure"] for t in biten), Decimal(0)) / len(biten), "0.1") if biten else None,
                          "hedefte": yuvarla(Decimal(sum(1 for t in biten if not t["asim"])) * 100 / len(biten)) if biten else None,
                          "kayip": sum(t["kayip_gece"] for t in l)})
        if sira:
            return sorted(satir, key=lambda s: sira.index(s["ad"]))
        return sorted(satir, key=lambda s: (-s["adet"], katla(s["ad"])))
    acik = sorted((t for t in talepler if t["acik"]), key=lambda t: (ONCELIK.index(t["oncelik"]), not t["asim"], t["bildirim"]))
    return {"talepler": sorted(talepler, key=lambda t: t["bildirim"]), "acik": acik, "kategori": ozet(lambda t: t["kategori"]), "oncelik": ozet(lambda t: t["oncelik"], list(ONCELIK)),
            "oda": ozet(lambda t: t["oda"]), "toplam": ozet(lambda t: "TOPLAM")[0], "tekrar": tekrarlar, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Açık Talepler"
    _baslik(ws, ["Talep No", "Oda", "Bildirim", "Kategori", "Açıklama", "Öncelik", "Durum", "Bekleme (saat)", "Hedef (saat)", "Satışa Kapalı", "Sorumlu"], (10, 7, 16, 16, 40, 8, 13, 10, 9, 8, 14))
    for t in r["acik"]:
        n = _satir(ws, [t["no"], t["oda"], t["bildirim"], t["kategori"], t["aciklama"], t["oncelik"], t["durum"], t["sure"], t["hedef"], "Evet" if t["kapali"] else "", t["sorumlu"]],
                   sar=(5,), dolgu={8: KIRMIZI if t["asim"] else None, 10: SARI if t["kapali"] else None})
        ws.cell(n, 3).number_format = "DD.MM.YYYY HH:MM"
    for ad, l in (("Kategori", r["kategori"]), ("Öncelik", r["oncelik"]), ("Oda", r["oda"])):
        ws = wb.create_sheet(ad)
        _baslik(ws, [ad, "Talep", "Açık", "Hedefi Geçen Açık", "Tamamlanan", "Ort. Çözüm (saat)", "Hedefte %", "Kayıp Oda-Gece"], (18, 8, 7, 10, 11, 11, 10, 10))
        for s in l + ([r["toplam"]] if ad == "Kategori" else []):
            _satir(ws, [s["ad"], s["adet"], s["acik"], s["geciken"], s["tamam"], s["ort"], s["hedefte"], s["kayip"]], para_s=(7,))
    ws = wb.create_sheet("Tekrar Eden")
    _baslik(ws, ["Oda", "Kategori", "Pencerede Talep", "Toplam Talep", "İlk", "Son"], (8, 18, 10, 10, 16, 16))
    for s in r["tekrar"]:
        n = _satir(ws, [s["oda"], s["kategori"], s["adet"], s["toplam"], s["ilk"], s["son"]])
        ws.cell(n, 5).number_format = ws.cell(n, 6).number_format = "DD.MM.YYYY HH:MM"
    ws = wb.create_sheet("Tüm Talepler")
    _baslik(ws, ["Talep No", "Oda", "Bildirim", "Kategori", "Öncelik", "Durum", "Tamamlanma", "Süre (saat)", "Hedef Aşıldı"], (10, 7, 16, 16, 8, 13, 16, 10, 8))
    for t in r["talepler"]:
        n = _satir(ws, [t["no"], t["oda"], t["bildirim"], t["kategori"], t["oncelik"], t["durum"], t["bitis"], t["sure"], "Evet" if t["asim"] else ""])
        ws.cell(n, 3).number_format = ws.cell(n, 7).number_format = "DD.MM.YYYY HH:MM"
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, simdi: datetime, hedef: dict | None = None, tekrar: int = 3, gun: int = 30) -> dict:
    talepler, uy = oku(yol)
    r = analiz(talepler, simdi, hedef or {"Acil": 4, "Normal": 24, "Düşük": 72}, tekrar, gun)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Oda arıza ve bakım taleplerini bekleme süresi, hedef süre ve tekrar eden arızalara göre takip eder.")
    p.add_argument("--talepler", type=Path, help="Talep No, Oda, Bildirim Zamanı, Kategori, Açıklama, Öncelik (Acil / Normal / Düşük), Durum, Tamamlanma Zamanı, Satışa Kapalı, Sorumlu")
    p.add_argument("--rapor-zamani", help="GG.AA.YYYY SS:DD (varsayılan şimdi; örnek veride 15.10.2026 09:00)")
    p.add_argument("--acil", type=int, default=4, help="Acil talepler için hedef süre, saat (varsayılan 4)")
    p.add_argument("--normal", type=int, default=24, help="Normal talepler için hedef süre, saat (varsayılan 24)")
    p.add_argument("--dusuk", type=int, default=72, help="Düşük öncelikli talepler için hedef süre, saat (varsayılan 72)")
    p.add_argument("--tekrar", type=int, default=3, help="Tekrar eden arıza için en az talep (varsayılan 3)")
    p.add_argument("--gun", type=int, default=30, help="Tekrar penceresi, gün (varsayılan 30)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "oda_ariza_takibi.xlsx")
    a = p.parse_args(argv)
    yol = a.talepler or ORNEK / "talepler.csv"
    simdi = zaman(a.rapor_zamani) if a.rapor_zamani else (datetime(2026, 10, 15, 9, 0) if a.talepler is None else datetime.now().replace(second=0, microsecond=0))
    if simdi is None or min(a.acil, a.normal, a.dusuk, a.gun) < 1 or a.tekrar < 2:
        print("[X] --rapor-zamani GG.AA.YYYY SS:DD; hedef süreler ve --gun en az 1, --tekrar en az 2 olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, simdi, {"Acil": a.acil, "Normal": a.normal, "Düşük": a.dusuk}, a.tekrar, a.gun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {t['adet']} talep · {t['acik']} açık ({t['geciken']} hedefi geçmiş) · ort. çözüm {t['ort']} saat · hedefte %{t['hedefte']} · kayıp oda-gece {t['kayip']}")
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
