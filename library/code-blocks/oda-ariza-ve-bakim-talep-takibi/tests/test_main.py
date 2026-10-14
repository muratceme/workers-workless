import sys
import tempfile
import unittest
from datetime import datetime
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
SIMDI = datetime(2026, 10, 15, 9, 0)
HEDEF = {"Acil": 4, "Normal": 24, "Düşük": 72}


def t(no, oda, b, kat, onc="Normal", bit=None, kapali=False):
    return {"no": no, "oda": oda, "bildirim": datetime(2026, 10, *b), "kategori": kat, "aciklama": "", "oncelik": onc, "durum": "Tamamlandı" if bit else "Açık",
            "bitis": datetime(2026, 10, *bit) if bit else None, "kapali": kapali, "sorumlu": ""}


def veri():
    return [t("1", "101", (14, 9), "Klima", "Acil"), t("2", "101", (1, 8), "Klima", bit=(1, 20)), t("3", "101", (10, 8), "Klima", "Düşük", bit=(14, 8)),
            t("4", "205", (12, 10), "Sıhhi", bit=(13, 16), kapali=True), t("5", "305", (13, 9), "Elektrik", "Düşük"), t("6", "305", (11, 9), "Sıhhi", "Normal", kapali=True)]


class AnalizTesti(unittest.TestCase):
    def test_sureler(self):
        r = main.analiz(veri(), SIMDI, HEDEF)
        s = {x["no"]: x for x in r["talepler"]}
        self.assertEqual([(s[n]["sure"], s[n]["asim"]) for n in "123456"],
                         [(D("24.0"), True), (D("12.0"), False), (D("96.0"), True), (D("30.0"), True), (D("48.0"), False), (D("96.0"), True)])
        self.assertEqual([s[n]["kayip_gece"] for n in ("4", "6")], [1, 4])
        self.assertEqual([x["no"] for x in r["acik"]], ["1", "6", "5"])
        tp = r["toplam"]
        self.assertEqual((tp["adet"], tp["acik"], tp["geciken"], tp["tamam"], tp["ort"], tp["hedefte"], tp["kayip"]), (6, 3, 2, 3, D("46.0"), D("33.33"), 5))
        self.assertEqual([(k["ad"], k["adet"]) for k in r["oncelik"]], [("Acil", 1), ("Normal", 3), ("Düşük", 2)])
        self.assertEqual(r["tekrar"][0]["adet"], 3)
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Orta", "Hedef süre aşıldı", "6 · Oda 305"), ("Orta", "Oda satışa kapalı", "6 · Oda 305"), ("Orta", "Tekrar eden arıza", "Oda 101"),
                          ("Yüksek", "Hedef süre aşıldı", "1 · Oda 101")])
        self.assertEqual(main.analiz(veri(), SIMDI, HEDEF, tekrar=3, gun=5)["tekrar"], [])
        with self.assertRaises(ValueError):
            main.analiz([], SIMDI, HEDEF)

    def test_donusumler(self):
        self.assertEqual([main._oncelik(x) for x in ("", "Kritik", "planlı", "orta", "x")], ["Normal", "Acil", "Düşük", "Normal", None])
        self.assertEqual([main._durum(x) for x in ("", "Parça bekliyor", "Çözüldü", "İptal edildi", "x")], ["Açık", "Devam ediyor", "Tamamlandı", "İptal", None])
        self.assertEqual([main.zaman(x) for x in ("15.10.2026 09:30", "2026-10-15 09:30:00", "15.10.2026", "x")],
                         [datetime(2026, 10, 15, 9, 30), datetime(2026, 10, 15, 9, 30), datetime(2026, 10, 15), None])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "talepler.csv", cikti, SIMDI)
            tp = r["toplam"]
            self.assertEqual((tp["adet"], tp["acik"], tp["geciken"], tp["ort"], tp["hedefte"], tp["kayip"]), (124, 4, 2, D("21.8"), D("79.17"), 9))
            self.assertEqual([u["kim"] for u in r["uyarilar"] if u["tur"] == "Hedef süre aşıldı"], ["T-0113 · Oda 302", "T-0077 · Oda 205"])
            self.assertIn(("Oda 214", "Klima"), [(u["kim"], u["aciklama"][:5]) for u in r["uyarilar"] if u["tur"] == "Tekrar eden arıza"])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Açık Talepler", "Kategori", "Öncelik", "Oda", "Tekrar Eden", "Tüm Talepler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "t.csv"
            y.write_text("Kayıt No;Oda No;Bildirim Tarihi;Arıza Türü;Öncelik;Durum;Kapanış;OOO\n1;101;14.10.2026 10:00;Klima;;;15.10.2026 08:00;\n1;101;14.10.2026 10:00;Klima;;;;\n"
                         "2;102;x;Klima;;;;\n3;103;14.10.2026 10:00;Klima;belki;;;\n4;104;14.10.2026 10:00;TV;;İptal;;\n5;105;14.10.2026 10:00;TV;;;13.10.2026 10:00;Evet\n"
                         "6;106;14.10.2026;TV;;Tamamlandı;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["durum"], x["kapali"]) for x in l], [u["tur"] for u in uy]),
                             ([("1", "Tamamlandı", False), ("5", "Açık", True), ("6", "Tamamlandı", False)],
                              ["Talep iki kez yazılmış", "Satır okunamadı", "Satır okunamadı", "Tamamlanma bildirimden önce", "Tamamlanma zamanı yok"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--talepler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--rapor-zamani", "15.10.2026 12:00", "--acil", "2"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--rapor-zamani", "x"], ["--normal", "0"], ["--tekrar", "1"])], [2] * 3)
            self.assertEqual(main.main(["--talepler", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Oda;Bildirim\n101;x\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
