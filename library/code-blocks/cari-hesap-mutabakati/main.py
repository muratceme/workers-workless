"""
Cari Hesap Mutabakatı — Workers / Workless kod bloğu
Muhasebe › Muhasebe Elemanı

Firmanın kendi cari ekstresi (ör. 320 Satıcılar / 120 Alıcılar muavini) ile karşı tarafın gönderdiği
ekstreyi karşılaştırır. Karşı tarafın kayıtları ters yönlüdür (bizim borcumuz onun alacağıdır); yön
otomatik algılanır. Eşleştirme sırası:
  1. Belge no + tutar      → mutabık
  2. Belge no, farklı tutar → tutar farkı (KDV / KDV tevkifatı ihtimali ayrıca denetlenir)
  3. Tutar + tarih (± gün)  → mutabık (belge nosu olmayan ödeme, havale, çek)
  4. Aynı tutar, uzak tarih → tarih farkı (bakiyeyi etkilemez)
Kalan kalemler belge türüne (fatura, iade, ödeme, çek/senet, kur/vade farkı) ve dönem sonuna yakınlığına
göre olası nedenleriyle listelenir; bakiye farkının tamamı kalem kalem açıklanır. İnternete bağlanmaz.

Kullanım:
    python main.py                                              # örnek ekstrelerle dener
    python main.py --biz bizim_ekstre.xlsx --karsi karsi_taraf_ekstresi.xlsx
    python main.py --biz bizim.xlsx --karsi karsi.xlsx --baslangic 01.07.2026 --bitis 30.09.2026 --tolerans 5
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
KDV_ORANLARI = (Decimal("0.20"), Decimal("0.10"), Decimal("0.01"))
TEVKIFAT_PAYLARI = (2, 3, 4, 5, 7, 9, 10)          # KDV tevkifat oranları: 2/10 … 10/10


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def para(x) -> Decimal:
    if x in (None, ""):
        return Decimal(0)
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("₺", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return Decimal(0)


def tarih(x) -> date | None:
    if x in (None, ""):
        return None
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(str(x).strip(), f).date()
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ayirici = csv.Sniffer().sniff(metin.splitlines()[0], delimiters=",;\t").delimiter
        satirlar = list(csv.reader(metin.splitlines(), delimiter=ayirici))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


ALANLAR = {
    "tarih": ("tarih", "belge tarihi", "fatura tarihi", "işlem tarihi", "fiş tarihi"),
    "belge": ("belge no", "fatura no", "evrak no", "belge numarası", "fatura numarası", "ettn"),
    "tur": ("belge türü", "evrak türü", "işlem türü", "fiş türü", "tür"),
    "aciklama": ("açıklama", "aciklama"),
    "borc": ("borç", "borc"),
    "alacak": ("alacak",),
    "tutar": ("tutar",),
}


# ----------------------------------------------------------------------------
# Kayıtlar
# ----------------------------------------------------------------------------

@dataclass
class Kayit:
    taraf: str              # "Biz" / "Karşı"
    satir: int
    tarih: date
    belge: str
    tur: str
    aciklama: str
    borc: Decimal
    alacak: Decimal
    tutar: Decimal = Decimal(0)      # bizim bakışımızla işaretli tutar: + karşı taraf bize borçlanır
    eslesme: str = ""
    not_: str = ""
    es: "Kayit | None" = field(default=None, repr=False)

    @property
    def anahtar(self) -> str:
        return belge_anahtari(self.belge)


def belge_anahtari(belge: str) -> str:
    return re.sub(r"[^0-9A-Z]", "", str(belge or "").upper().replace("İ", "I"))


def belge_ayni(a: str, b: str) -> bool:
    """Aynı belge mi? Tam eşitlik ya da (e-Fatura no ↔ kısa sıra no) rakam kuyruğu eşitliği."""
    if not a or not b:
        return False
    if a == b:
        return True
    ra, rb = re.sub(r"\D", "", a).lstrip("0"), re.sub(r"\D", "", b).lstrip("0")
    if not a.isdigit() and not b.isdigit():      # iki tam e-Fatura no farklıysa (ör. farklı seri) aynı belge değildir
        return False
    if a.isdigit() and b.isdigit():
        return ra == rb                           # 000123 = 123
    kisa, uzun = sorted((ra, rb), key=len)
    return len(kisa) >= 3 and uzun.endswith(kisa)


TURLER = [  # (tür, düzenli ifadeler) — sıra önemli; kelime başında aranır
    ("Devir", (r"devir", r"açılış")),
    ("İade", (r"iade",)),
    ("Kur farkı", (r"kur fark",)),
    ("Vade farkı", (r"vade fark",)),
    ("Çek / senet", (r"çek\b", r"cek\b", r"çeki", r"senet", r"bono")),
    ("Ödeme / tahsilat", (r"havale", r"eft\b", r"fast\b", r"ödeme", r"tahsilat", r"virman", r"kredi kartı", r"pos\b",
                          r"nakit", r"kasa\b", r"mahsup")),
    ("Fatura", (r"fatura", r"fat\.", r"ftr\b")),
]
_TUR_DESEN = [(ad, re.compile(r"\b(?:" + "|".join(ifadeler) + ")")) for ad, ifadeler in TURLER]


def tur_bul(tur: str, aciklama: str, belge: str) -> str:
    metin = f"{kucuk(tur)} {kucuk(aciklama)}"
    for ad, desen in _TUR_DESEN:
        if desen.search(metin):
            return ad
    if re.fullmatch(r"[A-Z0-9]{3}20\d{11}", belge_anahtari(belge)):     # e-Fatura / e-Arşiv numarası (16 hane)
        return "Fatura"
    return "Diğer"


def kayitlari_oku(yol: Path, taraf: str) -> list[Kayit]:
    satirlar = tablo_oku(yol)
    b = [kucuk(x) for x in satirlar[0]]
    k = {alan: next((i for i, x in enumerate(b) if x in adlar), None) for alan, adlar in ALANLAR.items()}
    if k["tarih"] is None or ((k["borc"] is None or k["alacak"] is None) and k["tutar"] is None):
        raise SystemExit(f"{yol.name}: Tarih ve Borç + Alacak (ya da işaretli Tutar) sütunları gerekli. Başlıklar: {satirlar[0]}")
    al = lambda r, a: r[k[a]] if k[a] is not None and k[a] < len(r) else None  # noqa: E731
    kayitlar = []
    for n, r in enumerate(satirlar[1:], 2):
        t = tarih(al(r, "tarih"))
        if t is None:
            continue
        if k["borc"] is not None and k["alacak"] is not None:
            borc, alacak = para(al(r, "borc")), para(al(r, "alacak"))
        else:
            v = para(al(r, "tutar"))
            borc, alacak = (v, Decimal(0)) if v >= 0 else (Decimal(0), -v)
        if borc == 0 and alacak == 0:
            continue
        belge, aciklama = str(al(r, "belge") or "").strip(), str(al(r, "aciklama") or "").strip()
        kayitlar.append(Kayit(taraf, n, t, belge, tur_bul(str(al(r, "tur") or ""), aciklama, belge), aciklama, borc, alacak))
    return kayitlar


def yon_belirle(biz: list[Kayit], karsi: list[Kayit]) -> bool:
    """Karşı tarafın kayıtları ters yönlü mü? Ortak belge numaralarında hangi yön tutarları tutturuyorsa o."""
    bizim = {(x.anahtar, x.borc - x.alacak) for x in biz if x.anahtar}
    ters = sum((y.anahtar, y.alacak - y.borc) in bizim for y in karsi if y.anahtar)
    duz = sum((y.anahtar, y.borc - y.alacak) in bizim for y in karsi if y.anahtar)
    if ters == duz == 0:   # belge no ortak değil: tutar kümelerine bak
        bizim_t = Counter(x.borc - x.alacak for x in biz)
        ters = sum(bizim_t[y.alacak - y.borc] > 0 for y in karsi)
        duz = sum(bizim_t[y.borc - y.alacak] > 0 for y in karsi)
    return ters >= duz


# ----------------------------------------------------------------------------
# Fark nedeni tahmini
# ----------------------------------------------------------------------------

def tutar_farki_nedeni(a: Decimal, b: Decimal) -> str:
    """Aynı belgede iki farklı tutar: KDV veya KDV tevkifatı farkı mı?"""
    kucuk_t, buyuk_t = sorted((abs(a), abs(b)))
    fark = buyuk_t - kucuk_t
    if fark <= Decimal("1"):
        return "Kuruş / yuvarlama farkı"
    pay_tol = Decimal("0.10")          # satır bazında KDV yuvarlamasına izin
    for oran in KDV_ORANLARI:
        if abs(kucuk_t * oran - fark) <= pay_tol:
            return f"Fark, küçük tutarın %{oran * 100:.0f}'si: bir taraf KDV hariç, diğeri KDV dahil kaydetmiş olabilir"
    for oran in (Decimal("0.20"), Decimal("0.10")):
        kdv = buyuk_t / (1 + oran) * oran
        for pay in TEVKIFAT_PAYLARI:
            if abs(kdv * pay / 10 - fark) <= pay_tol:
                return f"Fark, %{oran * 100:.0f} KDV'nin {pay}/10'una eşit: KDV tevkifatı bir tarafta düşülmemiş olabilir"
    return "Tutar farkı: fatura satırları, iskonto veya kur (dövizli fatura) karşılaştırılmalı"


def acik_kalem_nedeni(k: Kayit, donem_sonu: date, son_gun: int) -> str:
    kimde, kimde_yok = ("bizde", "karşı tarafta") if k.taraf == "Biz" else ("karşı tarafta", "bizde")
    yakin = (donem_sonu - k.tarih).days < son_gun
    neden = {
        "Fatura": f"Fatura {kimde_yok} kayıtlı değil: ulaşmamış, reddedilmiş/iptal edilmiş ya da başka cariye kaydedilmiş olabilir",
        "İade": f"İade faturası {kimde_yok} kayıtlı değil: kabul edilmemiş veya henüz işlenmemiş olabilir",
        "Ödeme / tahsilat": f"Ödeme {kimde_yok} görünmüyor: başka cariye / yanlış hesaba kaydedilmiş olabilir",
        "Çek / senet": f"Çek/senet {kimde_yok} kayıtlı değil: henüz ulaşmamış veya vade tarihiyle kaydedilmiş olabilir",
        "Kur farkı": f"Kur farkı faturası {kimde_yok} kayıtlı değil",
        "Vade farkı": f"Vade farkı faturası {kimde_yok} kayıtlı değil",
        "Devir": "Açılış (devir) bakiyesi farklı: önceki dönem mutabakatı kontrol edilmeli",
    }.get(k.tur, f"Kayıt yalnızca {kimde} var")
    if yakin and k.tur in ("Fatura", "İade", "Ödeme / tahsilat", "Çek / senet", "Kur farkı", "Vade farkı"):
        neden = f"Dönem sonuna {(donem_sonu - k.tarih).days} gün kala: yolda olabilir, {kimde_yok} sonraki döneme kaydedilmiş olabilir. " + neden
    return neden


# ----------------------------------------------------------------------------
# Eşleştirme
# ----------------------------------------------------------------------------

def eslestir(biz: list[Kayit], karsi: list[Kayit], tolerans: int) -> list[str]:
    uyarilar = []
    for liste in (biz, karsi):
        sayac = Counter((x.anahtar, x.tutar) for x in liste if x.anahtar)
        for (anahtar, tutar), n in sayac.items():
            if n > 1:
                taraf = liste[0].taraf
                uyarilar.append(f"{taraf}: {anahtar} numaralı belge aynı tutarla {n} kez kayıtlı (mükerrer kayıt olabilir)")

    def bagla(a: Kayit, b: Kayit, durum: str, not_: str = ""):
        a.es, b.es, a.eslesme, b.eslesme, a.not_, b.not_ = b, a, durum, durum, not_, not_

    acik_karsi = lambda: [y for y in karsi if y.es is None]  # noqa: E731

    # 0) Devir satırları birbiriyle karşılaştırılır
    for x in [x for x in biz if x.tur == "Devir"]:
        y = next((y for y in acik_karsi() if y.tur == "Devir"), None)
        if y is not None:
            if x.tutar == y.tutar:
                bagla(x, y, "Mutabık", "Açılış bakiyesi")
            else:
                bagla(x, y, "Tutar farkı", "Açılış (devir) bakiyesi farklı: önceki dönem mutabakatı kontrol edilmeli")

    # 1) Belge no + tutar, 2) belge no farklı tutar
    for x in [x for x in biz if x.es is None and x.anahtar]:
        adaylar = [y for y in acik_karsi() if belge_ayni(x.anahtar, y.anahtar) and (y.tutar > 0) == (x.tutar > 0)]
        tam = next((y for y in adaylar if y.tutar == x.tutar), None)
        if tam is not None:
            bagla(x, tam, "Mutabık", "Belge no ve tutar")
        elif adaylar:
            y = min(adaylar, key=lambda y: abs(y.tutar - x.tutar))
            bagla(x, y, "Tutar farkı", tutar_farki_nedeni(x.tutar, y.tutar))

    # 3) Tutar + tarih (tolerans içinde, en yakın tarih)
    for x in sorted([x for x in biz if x.es is None], key=lambda x: x.tarih):
        adaylar = [y for y in acik_karsi() if y.tutar == x.tutar and abs((y.tarih - x.tarih).days) <= tolerans]
        if adaylar:
            y = min(adaylar, key=lambda y: (abs((y.tarih - x.tarih).days), y.satir))
            gun = abs((y.tarih - x.tarih).days)
            not_ = "Tutar ve tarih" + (f" ({gun} gün fark)" if gun else "")
            if x.anahtar and y.anahtar and not belge_ayni(x.anahtar, y.anahtar):
                not_ += "; belge numaraları farklı, kontrol edin"
            bagla(x, y, "Mutabık", not_)

    # 4) Aynı tutar, tolerans dışı tarih
    for x in sorted([x for x in biz if x.es is None and x.tur != "Devir"], key=lambda x: x.tarih):
        adaylar = [y for y in acik_karsi() if y.tutar == x.tutar and y.tur != "Devir"]
        if adaylar:
            y = min(adaylar, key=lambda y: abs((y.tarih - x.tarih).days))
            bagla(x, y, "Tarih farkı", f"Aynı tutar, {abs((y.tarih - x.tarih).days)} gün farklı tarihte kayıtlı: "
                                        "doğru kayıt mı ve aynı dönemde mi kontrol edin")
    return uyarilar


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
PARA = "#,##0.00"


def tl(x) -> str:
    return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def bakiye_metni(t: Decimal) -> str:
    if t == 0:
        return "0,00 TL (kapalı)"
    return f"{tl(abs(t))} TL " + ("karşı taraf bize borçlu" if t > 0 else "biz karşı tarafa borçluyuz")


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = f"A{ws.max_row + 1}"


def calistir(biz_yolu: Path, karsi_yolu: Path, cikti: Path, baslangic: date | None = None, bitis: date | None = None,
             tolerans: int = 5, son_gun: int = 5, biz_adi: str = "", karsi_adi: str = "") -> dict:
    biz, karsi = kayitlari_oku(biz_yolu, "Biz"), kayitlari_oku(karsi_yolu, "Karşı")
    if baslangic or bitis:
        icinde = lambda k: (not baslangic or k.tarih >= baslangic) and (not bitis or k.tarih <= bitis)  # noqa: E731
        biz, karsi = [k for k in biz if icinde(k)], [k for k in karsi if icinde(k)]
    if not biz or not karsi:
        raise SystemExit("Ekstrelerden biri (seçilen dönemde) boş.")
    ters = yon_belirle(biz, karsi)
    for x in biz:
        x.tutar = x.borc - x.alacak
    for y in karsi:
        y.tutar = (y.alacak - y.borc) if ters else (y.borc - y.alacak)
    uyarilar = eslestir(biz, karsi, tolerans)
    donem_sonu = bitis or max(k.tarih for k in biz + karsi)

    bizim_bakiye = sum((x.tutar for x in biz), Decimal(0))
    karsi_bakiye = sum((y.tutar for y in karsi), Decimal(0))
    fark = bizim_bakiye - karsi_bakiye

    farklar = []        # (tür, açıklama satırı, bakiyeye etkisi)
    for x in biz:
        if x.es is None:
            farklar.append(("Bizde var, karşıda yok", x, None, x.tutar, acik_kalem_nedeni(x, donem_sonu, son_gun)))
        elif x.eslesme == "Tutar farkı":
            farklar.append(("Tutar farkı", x, x.es, x.tutar - x.es.tutar, x.not_))
        elif x.eslesme == "Tarih farkı":
            farklar.append(("Tarih farkı", x, x.es, Decimal(0), x.not_))
    for y in karsi:
        if y.es is None:
            farklar.append(("Karşıda var, bizde yok", None, y, -y.tutar, acik_kalem_nedeni(y, donem_sonu, son_gun)))
    aciklanan = sum((f[3] for f in farklar), Decimal(0))
    ozet_tur = defaultdict(lambda: [0, Decimal(0)])
    for f in farklar:
        ozet_tur[f[0]][0] += 1
        ozet_tur[f[0]][1] += f[3]

    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    mutabik = fark == 0 and not any(f[0] != "Tarih farkı" for f in farklar)
    satirlar = [
        ["Cari hesap mutabakatı", f"{biz_adi or 'Biz'} ↔ {karsi_adi or 'Karşı taraf'}"],
        ["Dönem", f"{(baslangic or min(k.tarih for k in biz + karsi)):%d.%m.%Y} – {donem_sonu:%d.%m.%Y}"],
        ["Karşı taraf kayıt yönü", "Ters (bizim borcumuz onun alacağı) — otomatik algılandı" if ters else "Aynı yön — otomatik algılandı"],
        [],
        ["Bizim kayıtlara göre bakiye", float(bizim_bakiye), bakiye_metni(bizim_bakiye)],
        ["Karşı tarafın kayıtlarına göre bakiye", float(karsi_bakiye), "(bizim bakışımıza çevrilmiş) " + bakiye_metni(karsi_bakiye)],
        ["Fark (biz − karşı)", float(fark)],
        ["Kalemlerle açıklanan fark", float(aciklanan)],
        ["Açıklanamayan fark", float(fark - aciklanan)],
        ["Sonuç", "MUTABIK" if mutabik else ("BAKİYE TUTUYOR, kalem farkı var" if fark == 0 else "MUTABIK DEĞİL")],
        [],
        ["Fark türü", "Adet", "Bakiyeye etkisi"],
    ]
    for s in satirlar:
        ws.append(s)
    for r in (5, 6, 7, 8, 9):
        ws.cell(r, 2).number_format = PARA
    ws.cell(10, 2).fill = YESIL if mutabik else KIRMIZI
    ws.cell(1, 1).font = Font(bold=True, size=13)
    for h in ws[12]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for tur in ("Bizde var, karşıda yok", "Karşıda var, bizde yok", "Tutar farkı", "Tarih farkı"):
        if tur in ozet_tur:
            ws.append([tur, ozet_tur[tur][0], float(ozet_tur[tur][1])])
            ws.cell(ws.max_row, 3).number_format = PARA
    ws.append(["Eşleşen (mutabık) kalem", sum(1 for x in biz if x.eslesme == "Mutabık")])
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 70
    for i, u in enumerate(uyarilar):
        ws.cell(ws.max_row + 2 if i == 0 else ws.max_row + 1, 1, "Uyarı").font = Font(bold=True)
        ws.cell(ws.max_row, 2, u)

    d = wb.create_sheet("Farklar")
    _baslik(d, ["Fark Türü", "Belge Türü", "Bizde Tarih", "Bizde Belge No", "Bizde Tutar", "Karşıda Tarih", "Karşıda Belge No",
                "Karşıda Tutar", "Bakiyeye Etkisi", "Olası Neden / Yapılacak", "Açıklama"],
            (24, 16, 12, 20, 15, 12, 20, 15, 15, 90, 40))
    for tur, x, y, etki, neden in sorted(farklar, key=lambda f: ((f[1] or f[2]).tarih, f[0])):
        ana = x or y
        d.append([tur, ana.tur, x.tarih if x else None, x.belge if x else "", float(x.tutar) if x else None,
                  y.tarih if y else None, y.belge if y else "", float(y.tutar) if y else None, float(etki), neden, ana.aciklama])
        for c in (3, 6):
            d.cell(d.max_row, c).number_format = "DD.MM.YYYY"
        for c in (5, 8, 9):
            d.cell(d.max_row, c).number_format = PARA
        d.cell(d.max_row, 1).fill = SARI if tur == "Tarih farkı" else KIRMIZI
        d.cell(d.max_row, 10).alignment = Alignment(wrap_text=True, vertical="top")
    d.auto_filter.ref = d.dimensions

    e = wb.create_sheet("Eşleşenler")
    _baslik(e, ["Belge Türü", "Bizde Tarih", "Bizde Belge No", "Karşıda Tarih", "Karşıda Belge No", "Tutar", "Eşleşme Yolu", "Açıklama"],
            (16, 12, 20, 12, 20, 15, 50, 40))
    for x in sorted([x for x in biz if x.eslesme == "Mutabık"], key=lambda x: x.tarih):
        e.append([x.tur, x.tarih, x.belge, x.es.tarih, x.es.belge, float(x.tutar), x.not_, x.aciklama])
        for c in (2, 4):
            e.cell(e.max_row, c).number_format = "DD.MM.YYYY"
        e.cell(e.max_row, 6).number_format = PARA

    for ad, liste in (("Bizim Ekstre", biz), ("Karşı Taraf Ekstresi", karsi)):
        s = wb.create_sheet(ad)
        _baslik(s, ["Satır", "Tarih", "Belge No", "Belge Türü", "Açıklama", "Borç", "Alacak", "Bizim Bakışla Tutar", "Durum"],
                (7, 12, 20, 16, 40, 14, 14, 18, 14))
        for k in sorted(liste, key=lambda k: (k.tarih, k.satir)):
            s.append([k.satir, k.tarih, k.belge, k.tur, k.aciklama, float(k.borc), float(k.alacak), float(k.tutar), k.eslesme or "Açık"])
            s.cell(s.max_row, 2).number_format = "DD.MM.YYYY"
            for c in (6, 7, 8):
                s.cell(s.max_row, c).number_format = PARA
            if k.eslesme != "Mutabık":
                s.cell(s.max_row, 9).fill = KIRMIZI if k.eslesme != "Tarih farkı" else SARI

    m = wb.create_sheet("Mutabakat Mektubu")
    if bizim_bakiye > 0:
        bakiye_cumle = f"{tl(bizim_bakiye)} TL borç bakiyesi vermektedir (firmanız şirketimize borçludur)."
    elif bizim_bakiye < 0:
        bakiye_cumle = f"{tl(-bizim_bakiye)} TL alacak bakiyesi vermektedir (şirketimiz firmanıza borçludur)."
    else:
        bakiye_cumle = "bakiye vermemektedir (hesap kapalıdır)."
    metin = [
        "CARİ HESAP MUTABAKAT MEKTUBU",
        "",
        f"Sayın {karsi_adi or '[Karşı taraf unvanı]'},",
        "",
        f"Kayıtlarımıza göre {donem_sonu:%d.%m.%Y} tarihi itibarıyla şirketimiz nezdindeki cari hesabınız {bakiye_cumle}",
        "",
        "Bakiyenin kayıtlarınızla uyumlu olup olmadığını aşağıdaki bölümü doldurup imzalayarak bildirmenizi rica ederiz.",
        "Mutabık olmamanız durumunda hesap ekstrenizi göndermenizi rica ederiz.",
        "",
        f"Saygılarımızla,  {biz_adi or '[Şirket unvanı]'}",
        "",
        "[ ] Mutabıkız      [ ] Mutabık değiliz — bizim kayıtlarımıza göre bakiye: ..................... TL",
        "",
        "Kaşe / İmza / Tarih:",
    ]
    for satir in metin:
        m.append([satir])
    m["A1"].font = Font(bold=True, size=13)
    m.column_dimensions["A"].width = 120

    b = wb.create_sheet("Bilgi")
    for s in [["Yöntem", "Sırayla: devir bakiyeleri → belge no + tutar → belge no (tutar farkı) → tutar + tarih (± tolerans) → aynı tutar (tarih farkı)"],
              ["Bakış", "Tüm tutarlar bizim bakışımızla işaretlidir: + karşı taraf bize borçlanır, − biz borçlanırız"],
              ["Tolerans", f"Tutar + tarih eşleşmesinde ± {tolerans} gün; dönem sonuna {son_gun} günden yakın kalemler 'yolda' sayılır"],
              ["Belge no", "Harf/rakam dışındaki karakterler atılır; e-Fatura no (ABC2026000000123) ile kısa sıra no (123) eşleşebilir"],
              ["Not", "Neden sütunu kural tabanlı bir tahmindir; kesin neden belgeler karşılaştırılarak doğrulanmalıdır"]]:
        b.append(s)
    b.column_dimensions["A"].width = 14
    b.column_dimensions["B"].width = 130
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    return {"biz": biz, "karsi": karsi, "ters": ters, "bizim_bakiye": bizim_bakiye, "karsi_bakiye": karsi_bakiye,
            "fark": fark, "aciklanan": aciklanan, "farklar": farklar, "uyarilar": uyarilar, "mutabik": mutabik}


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description="İki cari ekstreyi karşılaştırıp farkları ve olası nedenlerini raporlar.")
    ap.add_argument("--biz", type=Path, default=BURASI / "ornek_veri" / "bizim_ekstre.csv", help="Kendi cari ekstremiz (.xlsx/.csv)")
    ap.add_argument("--karsi", type=Path, default=BURASI / "ornek_veri" / "karsi_taraf_ekstresi.csv", help="Karşı tarafın ekstresi")
    ap.add_argument("--baslangic", help="Dönem başı (GG.AA.YYYY); verilmezse tüm kayıtlar")
    ap.add_argument("--bitis", help="Dönem sonu (GG.AA.YYYY); verilmezse en son kayıt tarihi")
    ap.add_argument("--tolerans", type=int, default=5, help="Tutar + tarih eşleşmesinde izin verilen gün farkı (varsayılan 5)")
    ap.add_argument("--son-gun", type=int, default=5, help="Dönem sonuna bu kadar günden yakın kalemler 'yolda' sayılır")
    ap.add_argument("--biz-adi", default="", help="Şirketimizin unvanı (mektup için)")
    ap.add_argument("--karsi-adi", default="", help="Karşı tarafın unvanı (mektup için)")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "cari_mutabakat.xlsx")
    a = ap.parse_args(argv)
    if a.biz == BURASI / "ornek_veri" / "bizim_ekstre.csv":
        a.biz_adi, a.karsi_adi = a.biz_adi or "Örnek Gıda A.Ş.", a.karsi_adi or "Kurgu Ambalaj A.Ş."
        a.bitis = a.bitis or "30.09.2026"
    s = calistir(a.biz, a.karsi, a.cikti, tarih(a.baslangic), tarih(a.bitis), a.tolerans, a.son_gun, a.biz_adi, a.karsi_adi)
    print(f"[OK] Biz {len(s['biz'])} kayıt, karşı taraf {len(s['karsi'])} kayıt · "
          f"karşı taraf yönü: {'ters' if s['ters'] else 'aynı'}")
    print(f"     Bizim bakiye : {bakiye_metni(s['bizim_bakiye'])}")
    print(f"     Karşı bakiye : {bakiye_metni(s['karsi_bakiye'])}")
    durum = "[OK] MUTABIK" if s["mutabik"] else f"[!] Fark {tl(s['fark'])} TL · {len(s['farklar'])} fark kalemi"
    print(f"{durum} · açıklanamayan fark {tl(s['fark'] - s['aciklanan'])} TL")
    for u in s["uyarilar"]:
        print(f"[!] {u}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
