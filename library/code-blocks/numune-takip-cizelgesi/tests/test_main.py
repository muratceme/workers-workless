import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 8)


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "n.xlsx"
        cls.s = main.calistir(ORNEK / "numuneler.csv", cls.cikti, ORNEK / "numune_sureleri.csv", ORNEK / "siparisler.csv", BUGUN)
        cls.a = {}
        for a in cls.s["aksiyonlar"]:
            cls.a.setdefault(a["tur"], []).append(a)
        cls.n = {n.no: n for n in cls.s["numuneler"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_eslestirme(self):
        self.assertEqual(main.tur_bul("Pre-Production"), "PP")
        self.assertEqual(main.tur_bul("Salesman sample"), "SMS")
        self.assertEqual(main.yanit_bul("Approved with comments"), "Yorumlu onay")
        self.assertEqual(main.yanit_bul("Rejected"), "Revize")
        self.assertEqual((self.n["N-013"].hedef, self.n["N-013"].hedef_kaynak), (date(2026, 10, 8), "Talep + 7 gün"))

    def test_durumlar(self):
        self.assertEqual({k: self.n[k].durum for k in ("N-004", "N-007", "N-011", "N-012", "N-016", "N-017", "N-003")},
                         {"N-004": "Yanıt gecikti", "N-007": "GECİKTİ", "N-011": "Yaklaşıyor", "N-012": "Hazırlanıyor", "N-016": "Müşteride",
                          "N-017": "İptal", "N-003": "Yorumlu onay"})
        self.assertEqual((self.n["N-004"].gecikme, self.n["N-007"].gecikme, self.n["N-010"].gecikme), (6, 3, 4))

    def test_aksiyonlar(self):
        self.assertEqual([a["onem"] for a in self.s["aksiyonlar"]][:3], ["Kritik", "Yüksek", "Yüksek"])
        self.assertIn("PO-501 planlanan kesim 15.10.2026 (7 gün kaldı)", self.a["Kesim yaklaşıyor, PP onayı yok"][0]["aciklama"])
        self.assertEqual([(a["model"], a["numune"]) for a in self.a["Revize numune açılmamış"]], [("SW-02", "PP R1")])
        self.assertIn("8 gündür yanıt yok", self.a["Müşteri yanıtı bekleniyor"][0]["aciklama"])
        self.assertEqual(self.a["Fit onayı olmadan PP"][0]["model"], "PN-04")
        self.assertIn("termin bugün", next(a for a in self.a["Termin yaklaşıyor"] if a["model"] == "PN-04")["aciklama"])
        self.assertIn("gönderim tarihinden (29.08.2026) önce", self.a["Veri hatası"][0]["aciklama"])
        self.assertIn("3 kez revize", self.a["Çok revizyon"][0]["aciklama"])
        self.assertEqual(self.a["AWB yok"][0]["no"], "N-016")
        self.assertNotIn("Gönderim gecikti", {a["tur"] for a in self.s["aksiyonlar"] if a["model"] == "JK-05"})

    def test_performans(self):
        pp = next(x for x in self.s["performans"]["tur"] if x["ad"] == "PP")
        self.assertEqual((pp["gonderilen"], pp["zamaninda"], pp["ort_gecikme"]), (3, 2, 6))
        m = {x["ad"]: x for x in self.s["performans"]["musteri"]}
        self.assertEqual(m["Kurgu Retail DE"]["en_uzun"], 14)

    def test_excel(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Aksiyonlar", "Model Durumu", "Numuneler", "Performans"])
        md = wb["Model Durumu"]
        self.assertEqual([c.value for c in md[1]], ["Model", "Müşteri", "Proto", "Fit", "Size Set", "PP", "SMS", "TOP"])
        satir = next(r for r in md.iter_rows(min_row=2) if r[0].value == "TS-01")
        self.assertEqual(satir[3].value, "R2 · Yorumlu onay 04.09")
        self.assertEqual(wb["Aksiyonlar"]["I1"].value, "Yapılan / Tarih")


class KuralTesti(unittest.TestCase):
    def test_kesim_gecmis_ve_pp_yok(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t / "n.csv").write_text("Model;Numune Türü;Talep Tarihi;İstenen Tarih\nX;Fit;01.09.2026;20.10.2026\n", encoding="utf-8")
            (t / "s.csv").write_text("Numune Türü;Kesimden Önce Onay (gün)\nPP;10\n", encoding="utf-8")
            (t / "p.csv").write_text("Sipariş No;Model;Planlanan Kesim Tarihi\nPO-1;X;05.10.2026\n", encoding="utf-8")
            s = main.calistir(t / "n.csv", t / "o.xlsx", t / "s.csv", t / "p.csv", BUGUN)
            a = [x for x in s["aksiyonlar"] if x["onem"] == "Kritik"]
            self.assertEqual(len(a), 1)
            self.assertIn("(3 gün önceydi)", a[0]["aciklama"])
            self.assertIn("PP numunesi hiç açılmamış", a[0]["aciklama"])

    def test_termin_yok_ve_belirsiz_yanit(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t / "n.csv").write_text("Model;Numune Türü;Gönderim Tarihi;Yanıt Tarihi;Müşteri Yanıtı\nX;Özel;;;\nY;Fit;01.10.2026;03.10.2026;bakılacak\n",
                                     encoding="utf-8")
            s = main.calistir(t / "n.csv", t / "o.xlsx", bugun=BUGUN)
            turler = {(a["tur"], a["model"]) for a in s["aksiyonlar"]}
            self.assertIn(("Termin yok", "X"), turler)
            self.assertIn(("Yanıt yorumlanamadı", "Y"), turler)

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--numuneler", str(Path(tmp) / "yok.csv")]), 1)
            self.assertEqual(main.main(["--bugun", "2026-13-40", "--cikti", str(Path(tmp) / "x.xlsx")]), 1)


if __name__ == "__main__":
    unittest.main()
