"""
Numune Takip Çizelgesi — Workers / Workless kod bloğu
Tekstil ve Konfeksiyon › Müşteri Temsilciliği (Merchandising) › Development Merchandiser

Numune taleplerini (proto, fit, size set, PP, SMS, TOP) aşama, gönderim tarihi ve müşteri yanıtlarıyla takip eder:
  - Gönderilmemiş numunelerde istenen tarih geçmiş mi, yaklaşıyor mu? (istenen tarih yoksa talep + hazırlık süresi)
  - Müşteriye gönderilmiş ama yanıt süresi aşılmış numuneler (hatırlatma listesi).
  - Revize / red yanıtı gelmiş ama yeni revizyonu açılmamış numuneler.
  - Planlanan kesim tarihi yaklaşırken onaylı PP numunesi olmayan siparişler.
  - Fit onayı olmadan açılmış PP numunesi; çok revizyon; yorumlu onaylarda üretime aktarılacak yorumlar.
  - Kayıt kontrolleri: AWB'siz gönderim, yanıt tarihi gönderimden önce, gönderim talepten önce.
Rapor: model × numune türü durum matrisi, aksiyon listesi, tüm numuneler, zamanında gönderim performansı.
İnternete bağlanmaz.

Kullanım:
    python main.py                                                   # örnek: 6 model, durum tarihi 08.10.2026
    python main.py --numuneler numuneler.xlsx --sureler numune_sureleri.csv --siparisler siparisler.xlsx --bugun 08.10.2026
"""
from __future__ import annotations

import argparse
import csv
import re
import statistics
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")

TURLER = {"Proto": ("proto", "prototype", "development", "gelistirme", "ilk numune"), "Fit": ("fit",),
          "Size Set": ("size set", "beden seti", "beden serisi", "size"), "PP": ("pp", "pre production", "preproduction", "uretim oncesi"),
          "SMS": ("sms", "salesman", "satis numunesi", "sales"), "TOP": ("top", "top of production", "production sample", "uretim numunesi"),
          "Sevkiyat": ("shipment", "sevkiyat", "sevk numunesi")}
TUR_SIRA = ["Proto", "Fit", "Size Set", "PP", "SMS", "TOP", "Sevkiyat"]
NUMUNE_SUTUNLARI = {"no": ("numune no", "numune", "no", "sira no"), "model": ("model", "style", "model kodu", "artikel"),
                    "musteri": ("musteri", "buyer", "marka"), "sezon": ("sezon", "season"), "tur": ("numune turu", "tur", "asama", "type"),
                    "rev": ("revizyon", "rev", "versiyon"), "talep": ("talep tarihi", "talep", "acilis tarihi"),
                    "istenen": ("istenen tarih", "termin", "musteri termini", "deadline", "istenen"),
                    "gonderim": ("gonderim tarihi", "gonderim", "kargo tarihi", "sevk tarihi"), "awb": ("kargo awb no", "awb no", "awb", "kargo no", "takip no"),
                    "yanit_tarih": ("yanit tarihi", "musteri yanit tarihi", "donus tarihi"), "yanit": ("musteri yaniti", "yanit", "sonuc", "durum"),
                    "yorum": ("musteri yorumu", "yorum", "yorumlar", "comments"), "sorumlu": ("sorumlu", "temsilci", "merchandiser")}
SURE_SUTUNLARI = {"tur": ("numune turu", "tur"), "hazirlik": ("hazirlik suresi gun", "hazirlik suresi", "hazirlik"),
                  "yanit": ("musteri yanit suresi gun", "musteri yanit suresi", "yanit suresi"),
                  "kesim": ("kesimden once onay gun", "kesimden once onay", "kesim oncesi gun")}
SIPARIS_SUTUNLARI = {"no": ("siparis no", "po", "po no"), "model": ("model", "style"), "musteri": ("musteri", "buyer"),
                     "kesim": ("planlanan kesim tarihi", "kesim tarihi", "kesim"), "sevk": ("sevk tarihi", "ex factory", "termin")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tur_bul(ham) -> str:
    k = katla(ham)
    for tur, adlar in TURLER.items():
        if k in adlar or any(k.startswith(a + " ") for a in adlar):
            return tur
    return metin(ham) or "—"


def yanit_bul(ham) -> str:
    k = katla(ham)
    if not k:
        return ""
    if k.startswith(("iptal", "cancel")):
        return "İptal"
    if k.startswith(("red", "ret", "reject", "revize", "revise", "tekrar", "duzeltme")):
        return "Revize"
    if "yorum" in k or "comment" in k or "kosullu" in k:
        return "Yorumlu onay"
    if k.startswith(("onay", "approved", "ok", "kabul")):
        return "Onay"
    return "Belirsiz"


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Numune:
    satir: int
    no: str
    model: str
    musteri: str
    sezon: str
    tur: str
    rev: int
    talep: date | None
    istenen: date | None
    gonderim: date | None
    awb: str
    yanit_tarih: date | None
    yanit: str
    yanit_ham: str
    yorum: str
    sorumlu: str
    hedef: date | None = None
    hedef_kaynak: str = ""
    durum: str = ""
    gecikme: int | None = None
    notlar: list = field(default_factory=list)

    @property
    def etiket(self) -> str:
        return f"{self.tur} R{self.rev}"

    @property
    def onayli(self) -> bool:
        return self.yanit in ("Onay", "Yorumlu onay")


def sureleri_oku(yol: Path | None) -> dict[str, dict]:
    if yol is None:
        return {}
    sonuc = {}
    for r in kayitlar(yol, SURE_SUTUNLARI, ("tur",)):
        if r.get("tur"):
            sonuc[tur_bul(r["tur"])] = {a: int(float(str(r[a]).replace(",", "."))) if metin(r.get(a)) else None for a in ("hazirlik", "yanit", "kesim")}
    return sonuc


def siparisleri_oku(yol: Path | None) -> dict[str, list[dict]]:
    sonuc = defaultdict(list)
    if yol is None:
        return sonuc
    for r in kayitlar(yol, SIPARIS_SUTUNLARI, ("model",)):
        if r.get("model"):
            sonuc[metin(r["model"])].append({"no": metin(r.get("no")), "kesim": tarih(r.get("kesim")), "sevk": tarih(r.get("sevk"))})
    return sonuc


def numuneleri_oku(yol: Path) -> list[Numune]:
    sonuc = []
    for r in kayitlar(yol, NUMUNE_SUTUNLARI, ("model", "tur")):
        if not r.get("model") and not r.get("no"):
            continue
        rev = re.search(r"\d+", metin(r.get("rev")))
        sonuc.append(Numune(r["_satir"], metin(r.get("no")) or f"Satır {r['_satir']}", metin(r.get("model")), metin(r.get("musteri")), metin(r.get("sezon")),
                            tur_bul(r.get("tur")), int(rev.group()) if rev else 1, tarih(r.get("talep")), tarih(r.get("istenen")), tarih(r.get("gonderim")),
                            metin(r.get("awb")), tarih(r.get("yanit_tarih")), yanit_bul(r.get("yanit")), metin(r.get("yanit")), metin(r.get("yorum")),
                            metin(r.get("sorumlu"))))
    return sonuc


# ----------------------------------------------------------------------------
# Analiz
# ----------------------------------------------------------------------------

SIRA = {"Kritik": 0, "Yüksek": 1, "Orta": 2, "Bilgi": 3}


def analiz_et(numuneler: list[Numune], sureler: dict, siparisler: dict, bugun: date, uyari_gun: int = 3, varsayilan_yanit: int = 7,
              cok_revizyon: int = 3) -> dict:
    aksiyonlar = []

    def aksiyon(onem, tur, n: Numune | None, aciklama, model=None):
        aksiyonlar.append({"onem": onem, "tur": tur, "no": n.no if n else "", "model": n.model if n else model, "numune": n.etiket if n else "",
                           "musteri": n.musteri if n else "", "sorumlu": n.sorumlu if n else "", "aciklama": aciklama})
        if n:
            n.notlar.append(tur)

    gruplar = defaultdict(list)
    for n in numuneler:
        gruplar[(n.model, n.tur)].append(n)
    for lst in gruplar.values():
        lst.sort(key=lambda n: (n.rev, n.talep or date.min))

    for n in numuneler:
        s = sureler.get(n.tur, {})
        if n.istenen:
            n.hedef, n.hedef_kaynak = n.istenen, "İstenen tarih"
        elif n.talep and s.get("hazirlik"):
            n.hedef, n.hedef_kaynak = n.talep + timedelta(days=s["hazirlik"]), f"Talep + {s['hazirlik']} gün"
        # Kayıt tutarlılığı
        if n.gonderim and n.talep and n.gonderim < n.talep:
            aksiyon("Orta", "Veri hatası", n, f"Gönderim tarihi ({n.gonderim:%d.%m.%Y}) talep tarihinden ({n.talep:%d.%m.%Y}) önce")
        if n.yanit_tarih and n.gonderim and n.yanit_tarih < n.gonderim:
            aksiyon("Orta", "Veri hatası", n, f"Yanıt tarihi ({n.yanit_tarih:%d.%m.%Y}) gönderim tarihinden ({n.gonderim:%d.%m.%Y}) önce")
        if n.yanit and n.yanit != "İptal" and not n.gonderim:
            aksiyon("Orta", "Veri hatası", n, f"Müşteri yanıtı '{n.yanit_ham}' var ama gönderim tarihi boş")
        if n.yanit == "Belirsiz":
            aksiyon("Bilgi", "Yanıt yorumlanamadı", n, f"'{n.yanit_ham}' onay / revize / iptal olarak tanınmadı")
        # Durum
        if n.yanit == "İptal":
            n.durum = "İptal"
            continue
        if n.yanit in ("Onay", "Yorumlu onay", "Revize"):
            n.durum = {"Onay": "Onaylandı", "Yorumlu onay": "Yorumlu onay", "Revize": "Revize istendi"}[n.yanit]
        elif n.gonderim:
            gun = (bugun - n.gonderim).days
            sure = s.get("yanit") if s.get("yanit") is not None else (varsayilan_yanit if n.tur != "SMS" else None)
            n.durum = "Müşteride"
            if sure is not None and gun > sure:
                n.durum = "Yanıt gecikti"
                aksiyon("Orta", "Müşteri yanıtı bekleniyor", n, f"{n.gonderim:%d.%m.%Y} tarihinde gönderildi ({n.awb or 'AWB yok'}), {gun} gündür yanıt yok "
                        f"(beklenen {sure} gün). Müşteriye hatırlatma yapın.")
        else:
            if n.hedef is None:
                n.durum = "Hazırlanıyor"
                aksiyon("Bilgi", "Termin yok", n, "İstenen tarih ve hazırlık süresi yok; numune termini belirleyin")
            elif bugun > n.hedef:
                n.durum, n.gecikme = "GECİKTİ", (bugun - n.hedef).days
                aksiyon("Yüksek", "Gönderim gecikti", n, f"Termin {n.hedef:%d.%m.%Y} ({n.hedef_kaynak}), {n.gecikme} gün gecikti; hâlâ gönderilmedi")
            elif (n.hedef - bugun).days <= uyari_gun:
                n.durum = "Yaklaşıyor"
                aksiyon("Orta", "Termin yaklaşıyor", n, f"Termin {n.hedef:%d.%m.%Y} ({n.hedef_kaynak}); " + (f"{(n.hedef - bugun).days} gün kaldı" if n.hedef > bugun else "termin bugün"))
            else:
                n.durum = "Hazırlanıyor"
        if n.gonderim and n.hedef:
            n.gecikme = (n.gonderim - n.hedef).days
        if n.gonderim and not n.awb:
            aksiyon("Bilgi", "AWB yok", n, f"{n.gonderim:%d.%m.%Y} gönderimi için kargo / AWB numarası yazılmamış")
        if n.yanit == "Yorumlu onay" and n.yorum:
            aksiyon("Bilgi", "Yorumu aktar", n, f"Yorumlu onay: \"{n.yorum}\". Yorumun bir sonraki numuneye / üretime aktarıldığını teyit edin.")

    # Revizyon zinciri
    for (model, tur), lst in gruplar.items():
        aktif = [n for n in lst if n.durum != "İptal"]
        if aktif and aktif[-1].yanit == "Revize":
            n = aktif[-1]
            aksiyon("Yüksek", "Revize numune açılmamış", n, (f"{n.yanit_tarih:%d.%m.%Y} tarihinde " if n.yanit_tarih else "")
                    + f"revize istendi, {tur} R{n.rev + 1} açılmamış. Yorum: \"{n.yorum or '—'}\"")
        revize = sum(n.yanit == "Revize" for n in lst)
        if revize >= cok_revizyon:
            aksiyon("Bilgi", "Çok revizyon", lst[-1], f"{model} {tur}: {revize} kez revize istendi; tekrarlayan yorumlar için müşteriyle toplantı / "
                    "ölçü tablosu teyidi düşünün. Yorumlar: " + " | ".join(f"R{n.rev}: {n.yorum}" for n in lst if n.yanit == "Revize" and n.yorum))

    # Sıra: fit onayı olmadan PP
    modeller = sorted({n.model for n in numuneler})
    for model in modeller:
        fitler = [n for n in gruplar.get((model, "Fit"), []) if n.durum != "İptal"]
        pp = [n for n in gruplar.get((model, "PP"), []) if n.durum != "İptal"]
        if fitler and pp and not any(f.onayli for f in fitler):
            aksiyon("Orta", "Fit onayı olmadan PP", pp[0], f"{model}: PP numunesi açıldı ama fit numunesi henüz onaylanmadı "
                    f"({fitler[-1].etiket}: {fitler[-1].durum}). PP'nin geçerli ölçülerle dikildiğinden emin olun.")
        # Kesim öncesi PP onayı
        esik = sureler.get("PP", {}).get("kesim")
        for sp in siparisler.get(model, []):
            if not sp["kesim"] or esik is None:
                continue
            kalan = (sp["kesim"] - bugun).days
            if not any(n.onayli for n in pp) and kalan <= esik:
                son = pp[-1] if pp else None
                ne = f"son PP R{son.rev}: {son.durum}" if son else "PP numunesi hiç açılmamış"
                aksiyon("Kritik", "Kesim yaklaşıyor, PP onayı yok", son, (f"{sp['no']} planlanan kesim {sp['kesim']:%d.%m.%Y} "
                        + (f"({kalan} gün kaldı)" if kalan >= 0 else f"({-kalan} gün önceydi)") + f"; onaylı PP numunesi yok ({ne}). "
                        "Kesim tarihi veya PP onayı için müşteriyle görüşün."), model)

    # Performans
    gonderilen = [n for n in numuneler if n.gonderim and n.hedef]
    performans = {"tur": [], "sorumlu": [], "musteri": []}
    for anahtar, alan in (("tur", lambda n: n.tur), ("sorumlu", lambda n: n.sorumlu or "—")):
        g = defaultdict(list)
        for n in gonderilen:
            g[alan(n)].append(n)
        for ad in sorted(g, key=lambda a: (TUR_SIRA.index(a) if a in TUR_SIRA else 99, a)):
            lst = g[ad]
            gec = [n.gecikme for n in lst if n.gecikme > 0]
            performans[anahtar].append({"ad": ad, "gonderilen": len(lst), "zamaninda": len(lst) - len(gec), "oran": (len(lst) - len(gec)) / len(lst),
                                        "ort_gecikme": statistics.fmean(gec) if gec else 0})
    g = defaultdict(list)
    for n in numuneler:
        if n.gonderim and n.yanit_tarih and n.yanit_tarih >= n.gonderim:
            g[n.musteri or "—"].append((n.yanit_tarih - n.gonderim).days)
    for ad in sorted(g):
        performans["musteri"].append({"ad": ad, "yanit": len(g[ad]), "ort_gun": statistics.fmean(g[ad]), "en_uzun": max(g[ad])})

    aksiyonlar.sort(key=lambda a: (SIRA[a["onem"]], a["model"] or "", a["no"]))
    return {"aksiyonlar": aksiyonlar, "gruplar": gruplar, "modeller": modeller, "performans": performans}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Kritik": "FDE2E1", "Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE", "GECİKTİ": "FDE2E1", "Revize istendi": "FDE2E1",
        "Yanıt gecikti": "FFF4CE", "Yaklaşıyor": "FFF4CE", "Onaylandı": "E3F4E1", "Yorumlu onay": "E3F4E1", "Müşteride": "E8F0FE", "İptal": "EEEEEE"}
KONTROL = PatternFill("solid", fgColor="FFF4CE")
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def _boya(hucre, anahtar):
    if anahtar in RENK:
        hucre.fill = PatternFill("solid", fgColor=RENK[anahtar])


def rapor_yaz(cikti: Path, numuneler: list[Numune], s: dict, bugun: date) -> None:
    wb = Workbook()
    ak = wb.active
    ak.title = "Aksiyonlar"
    _baslik(ak, ["Önem", "Tür", "Model", "Numune", "Numune No", "Müşteri", "Sorumlu", "Açıklama", "Yapılan / Tarih"], (9, 26, 9, 11, 10, 18, 12, 90, 24))
    for a in s["aksiyonlar"]:
        ak.append([a["onem"], a["tur"], a["model"], a["numune"], a["no"], a["musteri"], a["sorumlu"], a["aciklama"], ""])
        _boya(ak.cell(ak.max_row, 1), a["onem"])
        ak.cell(ak.max_row, 8).alignment = UST
        ak.cell(ak.max_row, 9).fill = KONTROL
    ak.auto_filter.ref = f"A1:I{ak.max_row}"

    md = wb.create_sheet("Model Durumu")
    turler = [t for t in TUR_SIRA if any(k[1] == t for k in s["gruplar"])] + sorted({k[1] for k in s["gruplar"]} - set(TUR_SIRA))
    _baslik(md, ["Model", "Müşteri"] + turler, (10, 18) + (24,) * len(turler))
    for model in s["modeller"]:
        musteri = next((n.musteri for n in numuneler if n.model == model and n.musteri), "")
        md.append([model, musteri])
        for j, t in enumerate(turler, 3):
            lst = s["gruplar"].get((model, t))
            if not lst:
                continue
            n = lst[-1]
            ek = (f" {n.yanit_tarih:%d.%m}" if n.yanit_tarih and n.yanit else f" (gönd. {n.gonderim:%d.%m})" if n.gonderim
                  else f" (termin {n.hedef:%d.%m})" if n.hedef else "")
            c = md.cell(md.max_row, j, f"R{n.rev} · {n.durum}{ek}")
            _boya(c, n.durum)
    md.append([])
    md.append([f"Durum tarihi {bugun:%d.%m.%Y}. Hücreler her türün son revizyonunu gösterir."])

    tk = wb.create_sheet("Numuneler")
    _baslik(tk, ["Numune No", "Model", "Müşteri", "Sezon", "Tür", "Rev.", "Talep", "Termin", "Termin Kaynağı", "Gönderim", "AWB", "Yanıt Tarihi",
                 "Yanıt", "Müşteri Yorumu", "Sorumlu", "Durum", "Gecikme (gün)", "İşaretler"],
            (10, 9, 18, 7, 9, 5, 11, 11, 16, 11, 11, 11, 13, 50, 12, 15, 9, 34))
    for n in sorted(numuneler, key=lambda n: (n.model, TUR_SIRA.index(n.tur) if n.tur in TUR_SIRA else 99, n.rev)):
        tk.append([n.no, n.model, n.musteri, n.sezon, n.tur, n.rev, n.talep, n.hedef, n.hedef_kaynak, n.gonderim, n.awb, n.yanit_tarih, n.yanit_ham,
                   n.yorum, n.sorumlu, n.durum, n.gecikme, "; ".join(dict.fromkeys(n.notlar))])
        for j in (7, 8, 10, 12):
            tk.cell(tk.max_row, j).number_format = "DD.MM.YYYY"
        _boya(tk.cell(tk.max_row, 16), n.durum)
        if n.gecikme and n.gecikme > 0:
            tk.cell(tk.max_row, 17).font = Font(bold=True, color="C00000")
        tk.cell(tk.max_row, 14).alignment = UST
    tk.auto_filter.ref = f"A1:R{tk.max_row}"

    pf = wb.create_sheet("Performans")
    pf.append(["Zamanında gönderim (gönderim ≤ termin)"])
    pf.cell(pf.max_row, 1).font = Font(bold=True)
    for baslik, anahtar in (("Numune Türü", "tur"), ("Sorumlu", "sorumlu")):
        pf.append([baslik, "Gönderilen", "Zamanında", "Oran", "Geç Gönderimde Ort. Gecikme (gün)"])
        for h in pf[pf.max_row]:
            h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        for x in s["performans"][anahtar]:
            pf.append([x["ad"], x["gonderilen"], x["zamaninda"], x["oran"], round(x["ort_gecikme"], 1)])
            pf.cell(pf.max_row, 4).number_format = "0%"
        pf.append([])
    pf.append(["Müşteri", "Yanıtlanan", "Ort. Yanıt Süresi (gün)", "En Uzun (gün)"])
    for h in pf[pf.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for x in s["performans"]["musteri"]:
        pf.append([x["ad"], x["yanit"], round(x["ort_gun"], 1), x["en_uzun"]])
    for j, w in enumerate((22, 12, 12, 10, 20), 1):
        pf.column_dimensions[get_column_letter(j)].width = w
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(numune_yolu: Path, cikti: Path, sure_yolu: Path | None = None, siparis_yolu: Path | None = None, bugun: date | None = None,
             uyari_gun: int = 3, varsayilan_yanit: int = 7) -> dict:
    bugun = bugun or date.today()
    numuneler = numuneleri_oku(numune_yolu)
    sureler = sureleri_oku(sure_yolu)
    siparisler = siparisleri_oku(siparis_yolu)
    s = analiz_et(numuneler, sureler, siparisler, bugun, uyari_gun, varsayilan_yanit)
    rapor_yaz(cikti, numuneler, s, bugun)
    return {**s, "numuneler": numuneler, "sureler": sureler, "bugun": bugun}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Numune taleplerini aşama, gönderim ve müşteri yanıtıyla takip eder; geciken numuneleri ve aksiyonları listeler.")
    p.add_argument("--numuneler", type=Path, default=ORNEK / "numuneler.csv", help="Numune listesi (.xlsx/.csv)")
    p.add_argument("--sureler", type=Path, default=ORNEK / "numune_sureleri.csv", help="Tür bazında hazırlık / yanıt / kesim öncesi onay süreleri")
    p.add_argument("--siparisler", type=Path, help="İsteğe bağlı: Model, Sipariş No, Planlanan Kesim Tarihi")
    p.add_argument("--bugun", help="Durum tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 08.10.2026)")
    p.add_argument("--uyari-gun", type=int, default=3, help="Termine kaç gün kala 'yaklaşıyor' denir (varsayılan 3)")
    p.add_argument("--yanit-gun", type=int, default=7, help="Süre tablosunda yoksa beklenen müşteri yanıt süresi (varsayılan 7 gün)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "numune_takip.xlsx")
    a = p.parse_args(argv)
    ornek = a.numuneler == ORNEK / "numuneler.csv"
    if ornek and a.siparisler is None:
        a.siparisler = ORNEK / "siparisler.csv"
    for y in (a.numuneler, a.sureler, a.siparisler):
        if y is not None and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    bugun = tarih(a.bugun) if a.bugun else (date(2026, 10, 8) if ornek else None)
    if a.bugun and bugun is None:
        print("[X] --bugun GG.AA.YYYY biçiminde olmalı")
        return 1
    try:
        s = calistir(a.numuneler, a.cikti, a.sureler, a.siparisler, bugun, a.uyari_gun, a.yanit_gun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(s['numuneler'])} numune · {len(s['modeller'])} model · {len(s['aksiyonlar'])} aksiyon (durum tarihi {s['bugun']:%d.%m.%Y})")
    for x in s["aksiyonlar"]:
        if x["onem"] in ("Kritik", "Yüksek"):
            print(f"[{'X' if x['onem'] == 'Kritik' else '!'}] {x['model']} {x['numune']}: {x['tur']} — {x['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
