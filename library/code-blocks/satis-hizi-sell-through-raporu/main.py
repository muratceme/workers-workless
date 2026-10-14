"""
Satış Hızı (Sell-Through) Raporu — Workers / Workless kod bloğu
Ürün Planlama ve Alokasyon › Alokasyon Uzmanı
Ürün ve mağaza bazında satış oranı ve stok günü raporlar, indirime aday ürünleri belirler.
  Satış oranı (sell-through) % = satılan ÷ (satılan + eldeki stok) × 100
  Günlük satış hızı             = son 7 gün satışı ÷ 7 (yazılmışsa), yoksa satılan ÷ girişten bu yana geçen gün
  Stok günü                     = eldeki stok ÷ günlük satış hızı
  İndirime aday                 = --sezon-sonu verilmişse stok günü sezon sonuna kalan günden fazla;
                                  verilmemişse satış oranı --esik (%40) altında — iki durumda da en az --min-hafta (4) hafta satışta
  Hızlı satıyor                 = stok günü --az (7) günden az
  Transfer önerisi              = aynı üründe stok günü --az altında bir mağaza ile --fazla (60) üstünde bir mağaza varsa
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"kod": ("urun kodu", "model kodu", "stok kodu", "kod", "sku"), "ad": ("urun adi", "model adi", "urun", "aciklama"), "magaza": ("magaza", "magaza adi", "sube", "lokasyon"),
         "giris": ("giris tarihi", "ilk sevk tarihi", "satisa baslama tarihi", "ilk giris"), "gelen": ("gelen adet", "sevk adedi", "giris adedi", "toplam sevk"),
         "satis": ("satis adedi", "satilan adet", "satilan", "satis"), "stok": ("eldeki stok", "stok", "stok adedi", "mevcut stok"), "son7": ("son 7 gun satis", "son 7 gun", "son hafta satis", "haftalik satis")}
SIFIR = Decimal(0)
YEDI = Decimal(7)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("kod", "satis", "stok")):
        kod = metin(r["kod"])
        if not kod:
            continue
        magaza = metin(r["magaza"]) or "Tümü"
        d = {a: para(r[a]) if metin(r[a]) else None for a in ("gelen", "satis", "stok", "son7")}
        giris = tarih(r["giris"]) if metin(r["giris"]) else None
        if d["satis"] is None or d["stok"] is None or any(metin(r[a]) and (d[a] is None or d[a] < 0) for a in d) or (metin(r["giris"]) and giris is None):
            uy.append(uyari("Yüksek", "Satır okunamadı", f"{kod} · {magaza}", f"Satır {r['_satir']}: sayılardan ya da giriş tarihinden biri okunamadı. Satır alınmadı."))
            continue
        if (katla(kod), katla(magaza)) in gorulen:
            uy.append(uyari("Yüksek", "Satır iki kez yazılmış", f"{kod} · {magaza}", f"Satır {r['_satir']}: aynı ürün ve mağaza ikinci kez geçiyor. İkinci satır alınmadı."))
            continue
        gorulen.add((katla(kod), katla(magaza)))
        if d["gelen"] is not None and d["gelen"] != d["satis"] + d["stok"]:
            uy.append(uyari("Bilgi", "Gelen ile satış + stok farklı", f"{kod} · {magaza}", f"Gelen {d['gelen']:,.0f}, satış + stok {d['satis'] + d['stok']:,.0f}. Transfer, iade ya da sayım farkı olabilir."))
        l.append({"kod": kod, "ad": metin(r["ad"]), "magaza": magaza, "giris": giris, **d})
    return l, uy


def _olcu(satis, stok, hiz, gun_sayisi) -> dict:
    return {"satis": satis, "stok": stok, "st": yuvarla(satis * 100 / (satis + stok)) if satis + stok else None, "hiz": yuvarla(hiz) if hiz is not None else None,
            "stok_gun": yuvarla(stok / hiz) if hiz else None, "gun": gun_sayisi, "hafta": yuvarla(Decimal(gun_sayisi) / YEDI, "0.1") if gun_sayisi is not None else None}


def hesapla(satirlar: list[dict], gun: date, sezon_sonu: date | None = None, esik: Decimal = Decimal(40), min_hafta: Decimal = Decimal(4), az: Decimal = Decimal(7), fazla: Decimal = Decimal(60)) -> dict:
    if not satirlar:
        raise ValueError("Raporlanacak satır yok.")
    kalan = (sezon_sonu - gun).days if sezon_sonu else None
    uy, urunler, detay = [], {}, []
    for s in satirlar:
        gecen = (gun - s["giris"]).days + 1 if s["giris"] else None
        if gecen is not None and gecen < 1:
            uy.append(uyari("Orta", "Giriş tarihi rapor tarihinden sonra", f"{s['kod']} · {s['magaza']}", f"Giriş {s['giris']:%d.%m.%Y}; satış hızı hesaplanmadı."))
            gecen = None
        hiz = s["son7"] / YEDI if s["son7"] is not None else (s["satis"] / gecen if gecen else None)
        m = {"kod": s["kod"], "ad": s["ad"], "magaza": s["magaza"], "giris": s["giris"], "gelen": s["gelen"], "_hiz": hiz, **_olcu(s["satis"], s["stok"], hiz, gecen)}
        m["durum"] = "Hızlı satıyor" if m["stok_gun"] is not None and m["stok_gun"] < az else "Satış yok" if not hiz and s["stok"] else ""
        detay.append(m)
        urunler.setdefault(katla(s["kod"]), []).append(m)
    sonuc = []
    for k, l in urunler.items():
        satis, stok = sum((m["satis"] for m in l), SIFIR), sum((m["stok"] for m in l), SIFIR)
        hizlar = [m["_hiz"] for m in l if m["_hiz"] is not None]
        gunler = [m["gun"] for m in l if m["gun"] is not None]
        u = {"kod": l[0]["kod"], "ad": next((m["ad"] for m in l if m["ad"]), ""), "magaza_sayisi": len(l), **_olcu(satis, stok, sum(hizlar, SIFIR) if hizlar else None, max(gunler) if gunler else None)}
        u["sezon_kalan"] = None
        if kalan is not None and u["hiz"] is not None:
            u["sezon_kalan"] = max(yuvarla(stok - sum(hizlar, SIFIR) * kalan, "1"), SIFIR)
        erken = u["hafta"] is not None and u["hafta"] < min_hafta
        if kalan is not None:
            aday = stok > 0 and (u["stok_gun"] is None or u["stok_gun"] > kalan)
            neden = ("satış yok" if u["stok_gun"] is None else f"stok {u['stok_gun']} günde erir") + f", sezon sonuna {kalan} gün var"
        else:
            aday = u["st"] is not None and u["st"] < esik and stok > 0
            neden = f"satış oranı %{u['st']}, eşik %{esik}"
        u["oneri"] = "Değerlendirmek için erken" if erken and aday else "İndirime aday" if aday else "Hızlı satıyor" if u["stok_gun"] is not None and u["stok_gun"] < az else "Normal"
        u["neden"] = neden if aday else ""
        azlar = sorted((m for m in l if m["stok_gun"] is not None and m["stok_gun"] < az), key=lambda m: m["stok_gun"])
        fazlalar = sorted((m for m in l if m["stok"] > 0 and (m["stok_gun"] is None or m["stok_gun"] > fazla)), key=lambda m: -(m["stok_gun"] or Decimal(10 ** 9)))
        u["transfer"] = "; ".join(f"{f['magaza']} → {a['magaza']}" for f, a in zip(fazlalar, azlar))
        if u["oneri"] == "İndirime aday":
            uy.append(uyari("Orta", "İndirime aday", u["kod"], f"{u['ad'] or u['kod']}: {neden}." + (f" Sezon sonunda tahmini {u['sezon_kalan']:,.0f} adet kalır." if u["sezon_kalan"] else "")))
        if u["transfer"]:
            uy.append(uyari("Bilgi", "Transfer önerisi", u["kod"], f"Stok fazlası olan mağazadan stoğu azalan mağazaya: {u['transfer']}."))
        u["magazalar"] = sorted(l, key=lambda m: katla(m["magaza"]))
        sonuc.append(u)
    sira = {"İndirime aday": 0, "Değerlendirmek için erken": 1, "Hızlı satıyor": 2, "Normal": 3}
    sonuc.sort(key=lambda u: (sira[u["oneri"]], u["st"] if u["st"] is not None else Decimal(101), katla(u["kod"])))
    detay.sort(key=lambda m: (katla(m["kod"]), katla(m["magaza"])))
    toplam = _olcu(sum((u["satis"] for u in sonuc), SIFIR), sum((u["stok"] for u in sonuc), SIFIR), None, None)
    return {"urunler": sonuc, "detay": detay, "kalan": kalan, "toplam": toplam, "uyarilar": uy}


ONERI_RENK = {"İndirime aday": KIRMIZI, "Değerlendirmek için erken": MAVI, "Hızlı satıyor": YESIL, "Normal": None}


def rapor_yaz(cikti: Path, r: dict, gun: date, esik: Decimal, az: Decimal) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ürünler"
    _baslik(ws, ["Ürün Kodu", "Ürün Adı", "Öneri", "Neden", "Mağaza", "Satılan", "Eldeki Stok", "Satış Oranı %", "Günlük Hız", "Stok Günü", "Satışta Hafta", "Sezon Sonu Tahmini Kalan", "Transfer Önerisi"],
            (14, 26, 16, 36, 7, 9, 9, 9, 9, 9, 8, 11, 30), "C2")
    for u in r["urunler"]:
        _satir(ws, [u["kod"], u["ad"], u["oneri"], u["neden"], u["magaza_sayisi"], u["satis"], u["stok"], u["st"], u["hiz"], u["stok_gun"], u["hafta"], u["sezon_kalan"], u["transfer"]],
               para_s=(8, 9, 10), dolgu={3: ONERI_RENK[u["oneri"]], 8: KIRMIZI if u["st"] is not None and u["st"] < esik else None}, sar=(4, 13))
    ws.append([])
    t = r["toplam"]
    ws.append([f"Rapor tarihi {gun:%d.%m.%Y}" + (f"; sezon sonuna {r['kalan']} gün" if r["kalan"] is not None else "") + f". Toplam satış oranı %{t['st']}."])
    ws = wb.create_sheet("Mağaza Detayı")
    _baslik(ws, ["Ürün Kodu", "Mağaza", "Durum", "Giriş Tarihi", "Gelen", "Satılan", "Eldeki Stok", "Satış Oranı %", "Günlük Hız", "Stok Günü", "Satışta Hafta"], (14, 18, 14, 11, 8, 8, 9, 9, 9, 9, 8), "C2")
    for m in r["detay"]:
        _satir(ws, [m["kod"], m["magaza"], m["durum"], m["giris"], m["gelen"], m["satis"], m["stok"], m["st"], m["hiz"], m["stok_gun"], m["hafta"]], para_s=(8, 9, 10), tarih_s=(4,),
               dolgu={3: YESIL if m["durum"] == "Hızlı satıyor" else SARI if m["durum"] else None, 10: SARI if m["stok_gun"] is not None and m["stok_gun"] < az else None})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, gun: date, sezon_sonu: date | None = None, esik: Decimal = Decimal(40), min_hafta: Decimal = Decimal(4), az: Decimal = Decimal(7), fazla: Decimal = Decimal(60)) -> dict:
    satirlar, uy = oku(yol)
    r = hesapla(satirlar, gun, sezon_sonu, esik, min_hafta, az, fazla)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, gun, esik, az)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Ürün ve mağaza bazında satış oranı ve stok günü raporlar, indirime aday ürünleri belirler.")
    p.add_argument("--satislar", type=Path, help="Ürün Kodu, Ürün Adı, Mağaza, Giriş Tarihi, Gelen Adet, Satış Adedi, Eldeki Stok, Son 7 Gün Satış")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 30.09.2026)")
    p.add_argument("--sezon-sonu", help="Sezonun bittiği gün GG.AA.YYYY; verilirse stok günü sezon sonuyla karşılaştırılır")
    p.add_argument("--esik", default="40", help="Sezon sonu verilmezse satış oranı bu yüzdenin altındaki ürünler indirime aday (varsayılan 40)")
    p.add_argument("--min-hafta", default="4", help="Bu kadar haftadan kısa süredir satıştaki ürünlere öneri verilmez (varsayılan 4)")
    p.add_argument("--az", default="7", help="Stok günü bunun altındaysa hızlı satıyor sayılır (varsayılan 7)")
    p.add_argument("--fazla", default="60", help="Transfer için stok günü bunun üstündeki mağaza fazla sayılır (varsayılan 60)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "satis_hizi_raporu.xlsx")
    a = p.parse_args(argv)
    yol = a.satislar or ORNEK / "urun_magaza_satislari.csv"
    gun = tarih(a.tarih) if a.tarih else (date(2026, 9, 30) if a.satislar is None else date.today())
    if a.satislar is None and not a.tarih and not a.sezon_sonu:
        a.sezon_sonu = "15.11.2026"
    sezon = tarih(a.sezon_sonu) if a.sezon_sonu else None
    sayilar = [para(x) for x in (a.esik, a.min_hafta, a.az, a.fazla)]
    if gun is None or (a.sezon_sonu and (sezon is None or sezon < gun)) or any(x is None or x < 0 for x in sayilar) or sayilar[0] > 100:
        print("[X] Tarihler GG.AA.YYYY olmalı (sezon sonu rapor tarihinden önce olamaz); --esik 0-100 arası, diğer sayılar sıfır ya da üstü")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, gun, sezon, *sayilar)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    say = lambda o: sum(1 for u in r["urunler"] if u["oneri"] == o)   # noqa: E731
    print(f"[OK] {len(r['urunler'])} ürün · {len(r['detay'])} ürün-mağaza · toplam satış oranı %{r['toplam']['st']} · indirime aday {say('İndirime aday')} · hızlı satan {say('Hızlı satıyor')}")
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
