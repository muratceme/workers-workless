import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
GUN = date(2026, 9, 30)


def s(kod, magaza, giris, satis, stok, son7=None, gelen=None):
    return {"kod": kod, "ad": "", "magaza": magaza, "giris": giris, "gelen": None if gelen is None else D(gelen), "satis": D(satis), "stok": D(stok), "son7": None if son7 is None else D(son7)}


def veri():
    g = date(2026, 9, 1)
    return [s("P1", "M1", g, 60, 40, 14), s("P1", "M2", g, 20, 80), s("p1", "M3", g, 30, 5, 7), s("P2", "M1", date(2026, 9, 20), 2, 50), s("P3", "M1", None, 0, 10)]


class HesapTesti(unittest.TestCase):
    def test_esik_kurali(self):
        r = main.hesapla(veri(), GUN)
        self.assertEqual([(u["kod"], u["oneri"], u["satis"], u["stok"], u["st"], u["hiz"], u["stok_gun"], u["hafta"], u["transfer"]) for u in r["urunler"]],
                         [("P3", "İndirime aday", D(0), D(10), D("0.00"), None, None, None, ""),
                          ("P2", "Değerlendirmek için erken", D(2), D(50), D("3.85"), D("0.18"), D("275.00"), D("1.6"), ""),
                          ("P1", "Normal", D(110), D(125), D("46.81"), D("3.67"), D("34.09"), D("4.3"), "M2 → M3")])
        self.assertEqual([(m["kod"], m["magaza"], m["st"], m["hiz"], m["stok_gun"], m["durum"]) for m in r["detay"]],
                         [("P1", "M1", D("60.00"), D("2.00"), D("20.00"), ""), ("P1", "M2", D("20.00"), D("0.67"), D("120.00"), ""), ("p1", "M3", D("85.71"), D("1.00"), D("5.00"), "Hızlı satıyor"),
                          ("P2", "M1", D("3.85"), D("0.18"), D("275.00"), ""), ("P3", "M1", D("0.00"), None, None, "Satış yok")])
        self.assertEqual(sorted((u["tur"], u["kim"], u["aciklama"][-14:]) for u in r["uyarilar"]), [("Transfer önerisi", "P1", "zaya: M2 → M3."), ("İndirime aday", "P3", ".00, eşik %40.")])

    def test_sezon_sonu(self):
        r = main.hesapla(veri(), GUN, date(2026, 10, 31))
        self.assertEqual([(u["kod"], u["oneri"], u["sezon_kalan"], u["neden"]) for u in r["urunler"]],
                         [("P3", "İndirime aday", None, "satış yok, sezon sonuna 31 gün var"), ("P1", "İndirime aday", D(11), "stok 34.09 günde erir, sezon sonuna 31 gün var"),
                          ("P2", "Değerlendirmek için erken", D(44), "stok 275.00 günde erir, sezon sonuna 31 gün var")])
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]), [("Transfer önerisi", "P1"), ("İndirime aday", "P1"), ("İndirime aday", "P3")])
        r = main.hesapla(veri(), GUN, esik=D(50), min_hafta=D(0), az=D(30), fazla=D(300))
        self.assertEqual([(u["kod"], u["oneri"]) for u in r["urunler"]], [("P3", "İndirime aday"), ("P2", "İndirime aday"), ("P1", "İndirime aday")])
        self.assertEqual(main.hesapla([s("X", "M", date(2026, 10, 1), 1, 1)], GUN)["uyarilar"][0]["tur"], "Giriş tarihi rapor tarihinden sonra")
        with self.assertRaises(ValueError):
            main.hesapla([], GUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "urun_magaza_satislari.csv", cikti, GUN, date(2026, 11, 15))
            self.assertEqual([(u["kod"], u["oneri"], u["st"], u["stok_gun"], u["transfer"]) for u in r["urunler"]],
                             [("GM-2040", "İndirime aday", D("26.09"), D("257.83"), ""), ("MN-5100", "Değerlendirmek için erken", D("7.50"), D("129.50"), ""),
                              ("PN-3300", "Normal", D("41.50"), D("37.23"), ""), ("KZ-1001", "Normal", D("58.33"), D("30.17"), ""), ("TS-0800", "Normal", D("88.93"), D("22.35"), "Sahil → AVM")])
            self.assertEqual((r["urunler"][0]["sezon_kalan"], r["toplam"]["st"]), (D(140), D("63.53")))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Ürünler", "Mağaza Detayı", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "s.csv"
            y.write_text("SKU;Şube;Satılan;Stok;Gelen Adet;İlk Sevk Tarihi\nA;M;1.200;300;1.500;01.09.2026\na;m;1;1;;\nB;M;x;1;;\nC;M;1;-1;;\nD;M;1;1;;yok\nE;;5;5;9;\n;M;1;1;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["kod"], x["magaza"], x["satis"], x["stok"], x["gelen"], x["giris"]) for x in l], [x["tur"] for x in uy]),
                             ([("A", "M", D(1200), D(300), D(1500), date(2026, 9, 1)), ("E", "Tümü", D(5), D(5), D(9), None)],
                              ["Satır iki kez yazılmış", "Satır okunamadı", "Satır okunamadı", "Satır okunamadı", "Gelen ile satış + stok farklı"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--satislar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "30.09.2026", "--sezon-sonu", "31.12.2026", "--esik", "30", "--min-hafta", "2", "--az", "5", "--fazla", "90"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--tarih", "30.09.2026", "--sezon-sonu", "01.09.2026"], ["--esik", "101"], ["--az", "-1"])], [2] * 4)
            self.assertEqual(main.main(["--satislar", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Kod;Stok\nA;5\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
