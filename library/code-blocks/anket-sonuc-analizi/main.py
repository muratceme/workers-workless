"""
Anket Sonuç Analizi — Workers / Workless kod bloğu
Pazarlama › Pazarlama Uzmanı

Pazar araştırması anketlerinin kapalı ve açık uçlu cevaplarını analiz eder:
  - Soru türleri (soru listesinden): Tek seçim, Çoklu seçim, Ölçek (ör. 1–5), Sayı, Açık uçlu, Demografi.
  - Tek / çoklu seçim: frekans, yüzde ve %95 hata payı (±1,96 × √(p(1−p)/n)). Çoklu seçimde yüzde, soruyu yanıtlayanlara
    göredir (toplam %100'ü aşabilir).
  - Ölçek: ortalama, medyan, standart sapma, üst iki kutu (en yüksek iki puan) ve alt iki kutu oranı.
  - Çapraz tablo (--kirilim sütunları): sütun yüzdeleri; tek seçim sorularında ki-kare bağımsızlık testi (p değeri).
    Beklenen frekansı 5'in altında olan hücre oranı %20'yi aşarsa test güvenilmez olarak işaretlenir. Ölçek sorularında
    grup ortalamaları.
  - Açık uçlu: kod çerçevesindeki (Kod; Anahtar Kelimeler) anahtar kelimelerle kodlama; bir yanıt birden çok kod
    alabilir; kodlanamayan yanıtlar elle kodlama için ayrı listelenir. Örnek alıntılarda e-posta ve telefon maskelenir.
  - Veri kalitesi: mükerrer yanıt numarası, geçersiz değer, soru bazında boş oranı, tüm ölçek sorularına aynı puanı
    veren yanıtlar (düz çizgi).
Rapor: özet, frekanslar, ölçek soruları, çapraz tablolar, anlamlı farklar, açık uçlu kodlar, kodlanmayan yanıtlar,
uyarılar. İnternete bağlanmaz.

Kullanım:
    python main.py                                                  # örnek: 240 yanıt, 10 soru
    python main.py --yanitlar yanitlar.xlsx --sorular sorular.xlsx --kodlar kod_cercevesi.xlsx --kirilim "Yaş Grubu" Cinsiyet
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import statistics
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
TURLER = {"tek secim": "tek", "tek": "tek", "tek secimli": "tek", "coklu secim": "coklu", "coklu": "coklu", "coklu secimli": "coklu", "olcek": "olcek",
          "likert": "olcek", "puan": "olcek", "sayi": "sayi", "sayisal": "sayi", "acik uclu": "acik", "acik": "acik", "metin": "acik", "demografi": "demografi",
          "demografik": "demografi"}
TUR_ADI = {"tek": "Tek seçim", "coklu": "Çoklu seçim", "olcek": "Ölçek", "sayi": "Sayı", "acik": "Açık uçlu", "demografi": "Demografi"}
_MASKE = [(re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+"), "[E-POSTA]"),
          (re.compile(r"(?:\+?90[\s-]?)?\(?0?5\d{2}\)?[\s.-]?\d{3}[\s.-]?\d{2}[\s.-]?\d{2}"), "[TELEFON]")]


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def metin(x) -> str:
    if isinstance(x, float) and x.is_integer():
        return str(int(x))
    return str(x if x is not None else "").strip()


def sayi(x) -> float | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip().replace("TL", "").replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def maskele(s: str) -> str:
    for desen, yerine in _MASKE:
        s = desen.sub(yerine, s)
    return s


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


# ----------------------------------------------------------------------------
# İstatistik
# ----------------------------------------------------------------------------

def _gamma_q(a: float, x: float) -> float:
    """Düzenlenmiş üst tamamlanmamış gama fonksiyonu Q(a, x) (seri ve sürekli kesir açılımı)."""
    if x <= 0:
        return 1.0
    lg = math.lgamma(a)
    if x < a + 1:
        toplam = terim = 1.0 / a
        n = a
        for _ in range(500):
            n += 1
            terim *= x / n
            toplam += terim
            if abs(terim) < abs(toplam) * 1e-14:
                break
        return 1.0 - toplam * math.exp(-x + a * math.log(x) - lg)
    b = x + 1 - a
    c = 1e300
    d = 1 / b
    h = d
    for i in range(1, 500):
        an = -i * (i - a)
        b += 2
        d = an * d + b
        d = 1e-300 if abs(d) < 1e-300 else d
        c = b + an / c
        c = 1e-300 if abs(c) < 1e-300 else c
        d = 1 / d
        delta = d * c
        h *= delta
        if abs(delta - 1) < 1e-14:
            break
    return math.exp(-x + a * math.log(x) - lg) * h


def ki_kare_p(ki: float, sd: int) -> float:
    return _gamma_q(sd / 2, ki / 2) if sd > 0 else 1.0


def ki_kare(tablo: list[list[int]]) -> dict | None:
    """Bağımsızlık testi. tablo[satır][sütun] = gözlenen frekans. Boş satır / sütunlar atılır."""
    tablo = [r for r in tablo if sum(r) > 0]
    if not tablo:
        return None
    sutunlar = [j for j in range(len(tablo[0])) if sum(r[j] for r in tablo) > 0]
    tablo = [[r[j] for j in sutunlar] for r in tablo]
    if len(tablo) < 2 or len(sutunlar) < 2:
        return None
    n = sum(map(sum, tablo))
    st, kt = [sum(r) for r in tablo], [sum(r[j] for r in tablo) for j in range(len(sutunlar))]
    ki, kucuk = 0.0, 0
    for i, r in enumerate(tablo):
        for j, g in enumerate(r):
            b = st[i] * kt[j] / n
            ki += (g - b) ** 2 / b
            kucuk += b < 5
    sd = (len(tablo) - 1) * (len(sutunlar) - 1)
    hucre = len(tablo) * len(sutunlar)
    return {"ki": ki, "sd": sd, "p": ki_kare_p(ki, sd), "n": n, "kucuk_oran": kucuk / hucre, "guvenilir": kucuk / hucre <= 0.20,
            "cramer_v": math.sqrt(ki / (n * (min(len(tablo), len(sutunlar)) - 1)))}


def p_yaz(p: float) -> str:
    return "< 0,0001" if p < 0.0001 else f"{p:.4f}".replace(".", ",")


def tr(x: float, basamak: int = 2) -> str:
    return f"{x:.{basamak}f}".replace(".", ",")


def hata_payi(p: float, n: int) -> float | None:
    return 1.96 * math.sqrt(p * (1 - p) / n) if n else None


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Soru:
    kod: str
    metin: str
    tur: str
    secenekler: list[str] = field(default_factory=list)
    sutun: int = -1


def secenek_ayir(s: str) -> list[str]:
    return [p.strip() for p in (s.split("|") if "|" in s else s.split(",")) if p.strip()]


def sorulari_oku(yol: Path) -> list[Soru]:
    satirlar = tablo_oku(yol)
    b = [katla(c) for c in satirlar[0]]

    def al(r, *adlar):
        for a in adlar:
            if a in b and b.index(a) < len(r):
                return metin(r[b.index(a)])
        return ""
    sorular = []
    for r in satirlar[1:]:
        kod = al(r, "soru kodu", "kod", "sutun")
        tur = TURLER.get(katla(al(r, "tur", "soru turu")))
        if not kod:
            continue
        if not tur:
            raise ValueError(f"{yol.name}: '{kod}' sorusunun türü okunamadı ('{al(r, 'tur', 'soru turu')}'). Tek seçim, Çoklu seçim, Ölçek, Sayı, Açık uçlu veya Demografi yazın")
        sorular.append(Soru(kod, al(r, "soru", "soru metni") or kod, tur, secenek_ayir(al(r, "secenekler", "secenek"))))
    return sorular


def yanitlari_oku(yol: Path, sorular: list[Soru]) -> tuple[list[dict], list[dict]]:
    satirlar = tablo_oku(yol)
    b = [katla(c) for c in satirlar[0]]
    uy = []
    for s in sorular:
        if katla(s.kod) in b:
            s.sutun = b.index(katla(s.kod))
        else:
            uy.append({"onem": "Yüksek", "tur": "Sütun bulunamadı", "kim": s.kod, "aciklama": "Soru listesindeki kod yanıt dosyasında sütun olarak yok; soru analiz edilmedi"})
    no_sutun = next((b.index(a) for a in ("yanit no", "katilimci no", "id", "no") if a in b), None)
    yanitlar, gorulen = [], Counter()
    for i, r in enumerate(satirlar[1:], 2):
        no = metin(r[no_sutun]) if no_sutun is not None and no_sutun < len(r) else f"Y{i}"
        gorulen[no] += 1
        if gorulen[no] > 1:
            uy.append({"onem": "Orta", "tur": "Mükerrer yanıt no", "kim": no, "aciklama": f"Satır {i}: aynı yanıt numarası tekrar ediyor; satır alınmadı"})
            continue
        yanitlar.append({"no": no, "deger": {s.kod: (metin(r[s.sutun]) if 0 <= s.sutun < len(r) else "") for s in sorular}})
    return yanitlar, uy


def kodlari_oku(yol: Path | None) -> list[dict]:
    if not yol:
        return []
    satirlar = tablo_oku(yol)
    b = [katla(c) for c in satirlar[0]]
    kodlar = []
    for r in satirlar[1:]:
        al = lambda *adlar: next((metin(r[b.index(a)]) for a in adlar if a in b and b.index(a) < len(r)), "")  # noqa: E731
        kod, kelimeler = al("kod", "tema"), [katla(k) for k in re.split(r"[|,]", al("anahtar kelimeler", "kelimeler")) if katla(k)]
        if kod and kelimeler:
            kodlar.append({"soru": al("soru kodu", "soru") or "*", "kod": kod, "kelimeler": kelimeler})
    return kodlar


def acik_kodla(yanit: str, kodlar: list[dict], soru: str) -> list[str]:
    k = " " + katla(yanit) + " "
    sonuc = []
    for c in kodlar:
        if c["soru"] not in ("*", "") and katla(c["soru"]) != katla(soru):
            continue
        if any(" " + kel in k for kel in c["kelimeler"]) and c["kod"] not in sonuc:
            sonuc.append(c["kod"])
    return sonuc


# ----------------------------------------------------------------------------
# Analiz
# ----------------------------------------------------------------------------

def secim_coz(s: Soru, ham: str) -> tuple[list[str], list[str]]:
    """(geçerli seçenekler, geçersiz değerler). Seçenek listesi varsa yazım farkı (büyük-küçük harf) düzeltilir."""
    if not ham:
        return [], []
    parcalar = secenek_ayir(ham) if s.tur == "coklu" else [ham.strip()]
    if not s.secenekler:
        return parcalar, []
    harita = {katla(x): x for x in s.secenekler}
    gecerli = [harita[katla(p)] for p in parcalar if katla(p) in harita]
    return list(dict.fromkeys(gecerli)), [p for p in parcalar if katla(p) not in harita]


def analiz_et(sorular: list[Soru], yanitlar: list[dict], kodlar: list[dict], kirilimlar: list[str], min_grup: int = 30) -> dict:
    uy = []
    n = len(yanitlar)
    sorular = [s for s in sorular if s.sutun >= 0]
    kod_harita = {katla(s.kod): s for s in sorular}
    kirilim_sorulari = []
    for k in kirilimlar:
        s = kod_harita.get(katla(k))
        if not s:
            uy.append({"onem": "Orta", "tur": "Kırılım bulunamadı", "kim": k, "aciklama": "Soru listesinde bu kodda bir sütun yok"})
        else:
            kirilim_sorulari.append(s)

    cozulmus = defaultdict(dict)        # soru kodu → yanıt no → değer(ler)
    frekans, olcek, sayisal, acik, bos = [], [], [], {}, {}
    for s in sorular:
        dolu = 0
        gecersiz = Counter()
        if s.tur in ("tek", "coklu", "demografi"):
            say = Counter()
            for y in yanitlar:
                g, kotu = secim_coz(s, y["deger"][s.kod])
                gecersiz.update(kotu)
                if g:
                    dolu += 1
                    cozulmus[s.kod][y["no"]] = g
                    say.update(g)
            sira = s.secenekler or [k for k, _ in say.most_common()]
            for sec in sira:
                p = say[sec] / dolu if dolu else 0.0
                frekans.append({"soru": s, "secenek": sec, "n": say[sec], "yanitlayan": dolu, "oran": p, "hata": hata_payi(p, dolu)})
        elif s.tur in ("olcek", "sayi"):
            degerler = []
            alt, ust = (sayi(s.secenekler[0]), sayi(s.secenekler[-1])) if len(s.secenekler) >= 2 else (None, None)
            for y in yanitlar:
                ham = y["deger"][s.kod]
                v = sayi(ham)
                if not ham:
                    continue
                if v is None or (s.tur == "olcek" and alt is not None and not alt <= v <= ust):
                    gecersiz[ham] += 1
                    continue
                dolu += 1
                degerler.append(v)
                cozulmus[s.kod][y["no"]] = v
            if degerler:
                kayit = {"soru": s, "n": dolu, "ort": statistics.fmean(degerler), "medyan": statistics.median(degerler),
                         "sd": statistics.stdev(degerler) if len(degerler) > 1 else 0.0, "min": min(degerler), "max": max(degerler)}
                if s.tur == "olcek":
                    en_alt, en_ust = (alt, ust) if alt is not None else (min(degerler), max(degerler))
                    kayit.update({"ust2": sum(v >= en_ust - 1 for v in degerler) / dolu, "alt2": sum(v <= en_alt + 1 for v in degerler) / dolu,
                                  "dagilim": Counter(int(v) if float(v).is_integer() else v for v in degerler), "alt": en_alt, "ust": en_ust})
                    olcek.append(kayit)
                else:
                    sayisal.append(kayit)
        else:
            kod_say, ornek, kodsuz = Counter(), defaultdict(list), []
            for y in yanitlar:
                ham = y["deger"][s.kod]
                if not ham or katla(ham) in ("yok", "fikrim yok", ""):
                    continue
                dolu += 1
                bulunan = acik_kodla(ham, kodlar, s.kod)
                cozulmus[s.kod][y["no"]] = bulunan
                kod_say.update(bulunan)
                for c in bulunan:
                    if len(ornek[c]) < 3:
                        ornek[c].append(maskele(ham))
                if not bulunan:
                    kodsuz.append((y["no"], maskele(ham)))
            acik[s.kod] = {"soru": s, "n": dolu, "kodlar": kod_say, "ornek": ornek, "kodsuz": kodsuz}
            if dolu and len(kodsuz) / dolu > 0.25:
                uy.append({"onem": "Orta", "tur": "Kodlanmayan yanıt çok", "kim": s.kod,
                           "aciklama": f"{dolu} yanıtın {len(kodsuz)} tanesi hiçbir koda girmedi; kod çerçevesine anahtar kelime ekleyin"})
        bos[s.kod] = 1 - dolu / n if n else 0
        if gecersiz:
            uy.append({"onem": "Orta", "tur": "Geçersiz değer", "kim": s.kod,
                       "aciklama": f"{sum(gecersiz.values())} yanıt seçenek listesinde / ölçekte yok ve sayılmadı: " + ", ".join(f"'{k}' ({v})" for k, v in gecersiz.most_common(5))})
        if bos[s.kod] > 0.20 and s.tur != "acik":
            uy.append({"onem": "Bilgi", "tur": "Boş oranı yüksek", "kim": s.kod, "aciklama": f"Yanıtların %{bos[s.kod] * 100:.0f}'i boş"})

    # Düz çizgi: en az 4 ölçek sorusunun tamamına aynı puan
    olcek_kodlari = [s.kod for s in sorular if s.tur == "olcek"]
    duz = []
    if len(olcek_kodlari) >= 4:
        for y in yanitlar:
            v = [cozulmus[k].get(y["no"]) for k in olcek_kodlari]
            if None not in v and len(set(v)) == 1:
                duz.append(y["no"])
        if duz:
            uy.append({"onem": "Bilgi", "tur": "Düz çizgi yanıt", "kim": f"{len(duz)} yanıt",
                       "aciklama": f"Tüm ölçek sorularına aynı puan: {', '.join(duz[:15])}{'…' if len(duz) > 15 else ''}. Dikkatsiz yanıt olabilir; çıkarıp karşılaştırın"})

    # Çapraz tablolar
    capraz, anlamli = [], []
    for k in kirilim_sorulari:
        gruplar = k.secenekler or [g for g, _ in Counter(v[0] for v in cozulmus[k.kod].values()).most_common()]
        grup_n = Counter(v[0] for v in cozulmus[k.kod].values())
        for g in gruplar:
            if 0 < grup_n[g] < min_grup:
                uy.append({"onem": "Bilgi", "tur": "Küçük grup", "kim": f"{k.kod} = {g}", "aciklama": f"n = {grup_n[g]} (< {min_grup}); bu grubun yüzdeleri yön göstericidir"})
        for s in sorular:
            if s.kod == k.kod or s.tur in ("acik", "demografi", "sayi"):
                continue
            ortak = [no for no in cozulmus[s.kod] if no in cozulmus[k.kod]]
            if s.tur in ("tek", "coklu"):
                secenekler = s.secenekler or sorted({x for v in cozulmus[s.kod].values() for x in v})
                tablo = [[sum(1 for no in ortak if sec in cozulmus[s.kod][no] and cozulmus[k.kod][no][0] == g) for g in gruplar] for sec in secenekler]
                payda = [sum(1 for no in ortak if cozulmus[k.kod][no][0] == g) for g in gruplar]
                test = ki_kare(tablo) if s.tur == "tek" else None
                capraz.append({"soru": s, "kirilim": k, "gruplar": gruplar, "secenekler": secenekler, "tablo": tablo, "payda": payda, "test": test})
                if test and test["p"] < 0.05:
                    anlamli.append({"soru": s, "kirilim": k, "test": test,
                                    "ozet": _fark_ozeti(secenekler, gruplar, tablo, payda)})
            else:
                ort = []
                for g in gruplar:
                    v = [cozulmus[s.kod][no] for no in ortak if cozulmus[k.kod][no][0] == g]
                    ort.append((statistics.fmean(v) if v else None, len(v)))
                capraz.append({"soru": s, "kirilim": k, "gruplar": gruplar, "ortalamalar": ort})
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    uy.sort(key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))
    return {"n": n, "sorular": sorular, "frekans": frekans, "olcek": olcek, "sayisal": sayisal, "acik": acik, "bos": bos, "capraz": capraz, "anlamli": anlamli,
            "uyarilar": uy, "duz": duz, "hata_payi": hata_payi(0.5, n), "kirilimlar": kirilim_sorulari}


def _fark_ozeti(secenekler, gruplar, tablo, payda) -> str:
    """En büyük yüzde farkı olan seçeneği ve grupları yazar."""
    en = None
    for i, sec in enumerate(secenekler):
        oranlar = [(tablo[i][j] / payda[j], gruplar[j]) for j in range(len(gruplar)) if payda[j]]
        if len(oranlar) >= 2:
            buyuk, kucuk = max(oranlar), min(oranlar)
            if en is None or buyuk[0] - kucuk[0] > en[0]:
                en = (buyuk[0] - kucuk[0], sec, buyuk, kucuk)
    if not en:
        return ""
    return f"'{en[1]}': {en[2][1]} %{en[2][0] * 100:.0f}, {en[3][1]} %{en[3][0] * 100:.0f}"


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
UST = Alignment(vertical="top", wrap_text=True)
YF = "0.0%"


def _baslik(ws, basliklar, genislik=()):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def rapor_yaz(cikti: Path, s: dict) -> None:
    wb = Workbook()
    oz = wb.active
    oz.title = "Özet"
    _baslik(oz, ["Gösterge", "Değer"], (46, 40))
    for a, v in [("Yanıt sayısı (n)", s["n"]), ("Genel hata payı (%95 güven, p = 0,5)", f"± %{s['hata_payi'] * 100:.1f}".replace(".", ",") if s["hata_payi"] else "—"),
                 ("Soru sayısı", len(s["sorular"])), ("Kırılımlar", ", ".join(k.kod for k in s["kirilimlar"]) or "—"),
                 ("İstatistiksel olarak anlamlı fark (p < 0,05)", len(s["anlamli"])), ("Düz çizgi yanıt", len(s["duz"]))]:
        oz.append([a, v])
    oz.append([])
    _baslik(oz, ["Soru Kodu", "Soru", "Tür", "Boş Oranı"], ())
    oz.column_dimensions["C"].width, oz.column_dimensions["D"].width = 14, 10
    for q in s["sorular"]:
        oz.append([q.kod, q.metin, TUR_ADI[q.tur], s["bos"][q.kod]])
        oz.cell(oz.max_row, 4).number_format = "0%"
    oz.append([])
    oz.append(["Hata payı basit rastgele örneklem varsayar; kota veya kolayda örneklemde yalnız yön göstericidir."])

    fr = wb.create_sheet("Frekanslar")
    _baslik(fr, ["Soru Kodu", "Soru", "Tür", "Seçenek", "n", "Yanıtlayan", "Oran", "± Hata Payı"], (10, 46, 12, 30, 7, 10, 8, 10))
    for x in s["frekans"]:
        fr.append([x["soru"].kod, x["soru"].metin, TUR_ADI[x["soru"].tur], x["secenek"], x["n"], x["yanitlayan"], x["oran"], x["hata"]])
        fr.cell(fr.max_row, 7).number_format = fr.cell(fr.max_row, 8).number_format = YF
    fr.auto_filter.ref = f"A1:H{fr.max_row}"

    ol = wb.create_sheet("Ölçek Soruları")
    puanlar = sorted({p for x in s["olcek"] for p in x["dagilim"]})
    _baslik(ol, ["Soru Kodu", "Soru", "n", "Ortalama", "Medyan", "Std. Sapma", "Üst 2 Kutu", "Alt 2 Kutu"] + [f"{p} puan" for p in puanlar],
            [10, 46, 7, 9, 8, 9, 10, 10] + [8] * len(puanlar))
    for x in sorted(s["olcek"], key=lambda x: -x["ort"]):
        ol.append([x["soru"].kod, x["soru"].metin, x["n"], round(x["ort"], 2), x["medyan"], round(x["sd"], 2), x["ust2"], x["alt2"]]
                  + [x["dagilim"].get(p, 0) / x["n"] for p in puanlar])
        for j in range(7, 9 + len(puanlar)):
            ol.cell(ol.max_row, j).number_format = "0%"
    if s["sayisal"]:
        ol.append([])
        ol.append(["Sayı soruları"])
        ol.cell(ol.max_row, 1).font = Font(bold=True)
        ol.append(["Soru Kodu", "Soru", "n", "Ortalama", "Medyan", "Std. Sapma", "En Düşük", "En Yüksek"])
        for x in s["sayisal"]:
            ol.append([x["soru"].kod, x["soru"].metin, x["n"], round(x["ort"], 2), x["medyan"], round(x["sd"], 2), x["min"], x["max"]])

    cp = wb.create_sheet("Çapraz Tablolar")
    cp.column_dimensions["A"].width, cp.column_dimensions["B"].width = 12, 40
    for x in s["capraz"]:
        cp.append([f"{x['soru'].kod} × {x['kirilim'].kod}", x["soru"].metin])
        cp.cell(cp.max_row, 1).font = Font(bold=True)
        cp.append(["", "Seçenek / gösterge"] + x["gruplar"])
        for h in cp[cp.max_row]:
            h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        if "tablo" in x:
            for i, sec in enumerate(x["secenekler"]):
                cp.append(["", sec] + [x["tablo"][i][j] / x["payda"][j] if x["payda"][j] else None for j in range(len(x["gruplar"]))])
                for j in range(3, 3 + len(x["gruplar"])):
                    cp.cell(cp.max_row, j).number_format = "0%"
            cp.append(["", "n"] + x["payda"])
            t = x["test"]
            if t:
                cp.append(["", f"Ki-kare = {tr(t['ki'])}, sd = {t['sd']}, p = {p_yaz(t['p'])}, Cramér V = {tr(t['cramer_v'])}"
                           + ("" if t["guvenilir"] else " — beklenen frekansı 5'in altında çok hücre var; test güvenilmez")])
                cp.cell(cp.max_row, 2).fill = PatternFill("solid", fgColor="E3F4E1" if t["p"] < 0.05 and t["guvenilir"] else "FFFFFF")
        else:
            cp.append(["", "Ortalama"] + [None if o is None else round(o, 2) for o, _ in x["ortalamalar"]])
            cp.append(["", "n"] + [m for _, m in x["ortalamalar"]])
        cp.append([])

    an = wb.create_sheet("Anlamlı Farklar")
    _baslik(an, ["Soru", "Kırılım", "p", "Cramér V", "Test Güvenilir mi?", "En Büyük Fark", "Soru Metni"], (10, 14, 9, 9, 12, 60, 46))
    for x in sorted(s["anlamli"], key=lambda x: x["test"]["p"]):
        t = x["test"]
        an.append([x["soru"].kod, x["kirilim"].kod, p_yaz(t["p"]), round(t["cramer_v"], 2), "Evet" if t["guvenilir"] else "Hayır", x["ozet"], x["soru"].metin])

    ak = wb.create_sheet("Açık Uçlu Kodlar")
    _baslik(ak, ["Soru Kodu", "Kod (tema)", "Yanıt", "Yanıtlayanlara Oran", "Örnek Yanıtlar"], (10, 26, 8, 12, 110))
    for kod, x in s["acik"].items():
        for c, m in x["kodlar"].most_common():
            ak.append([kod, c, m, m / x["n"] if x["n"] else None, " | ".join(x["ornek"][c])])
            ak.cell(ak.max_row, 4).number_format = "0%"
            ak.cell(ak.max_row, 5).alignment = UST
        ak.append([kod, "(Kodlanmadı)", len(x["kodsuz"]), len(x["kodsuz"]) / x["n"] if x["n"] else None, "Kodlanmayan Yanıtlar sayfasına bakın"])
        ak.cell(ak.max_row, 4).number_format = "0%"

    kd = wb.create_sheet("Kodlanmayan Yanıtlar")
    _baslik(kd, ["Soru Kodu", "Yanıt No", "Yanıt", "Elle Kod"], (10, 10, 100, 24))
    for kod, x in s["acik"].items():
        for no, y in x["kodsuz"]:
            kd.append([kod, no, y, ""])
            kd.cell(kd.max_row, 3).alignment = UST
            kd.cell(kd.max_row, 4).fill = PatternFill("solid", fgColor="FFF4CE")

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Soru / Kim", "Açıklama"], (9, 24, 22, 110))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["kim"], u["aciklama"]])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(yanit_yolu: Path, soru_yolu: Path, cikti: Path, kod_yolu: Path | None = None, kirilimlar: list[str] | None = None, min_grup: int = 30) -> dict:
    sorular = sorulari_oku(soru_yolu)
    yanitlar, uy = yanitlari_oku(yanit_yolu, sorular)
    if not yanitlar:
        raise ValueError(f"{yanit_yolu.name}: yanıt bulunamadı")
    if kirilimlar is None:
        kirilimlar = [s.kod for s in sorular if s.tur == "demografi"]
    s = analiz_et(sorular, yanitlar, kodlari_oku(kod_yolu), kirilimlar, min_grup)
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    s["uyarilar"] = sorted(uy + s["uyarilar"], key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))
    if s["n"] < 100:
        s["uyarilar"].append({"onem": "Bilgi", "tur": "Küçük örneklem", "kim": f"n = {s['n']}", "aciklama": "100'ün altındaki örneklemde yüzdeler ve testler yön göstericidir"})
    rapor_yaz(cikti, s)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Anketlerin kapalı ve açık uçlu cevaplarını analiz eder: frekans, çapraz tablo, ki-kare, tema kodlama.")
    p.add_argument("--yanitlar", type=Path, default=ORNEK / "yanitlar.csv", help="Her satır bir yanıt: Yanıt No, demografi sütunları, soru kodu sütunları")
    p.add_argument("--sorular", type=Path, default=ORNEK / "sorular.csv",
                   help="Soru Kodu, Soru, Tür (Tek seçim / Çoklu seçim / Ölçek / Sayı / Açık uçlu / Demografi), Seçenekler ('|' ile; ölçekte 1|2|3|4|5)")
    p.add_argument("--kodlar", type=Path, help="Açık uçlu kod çerçevesi: Soru Kodu (veya *), Kod, Anahtar Kelimeler (',' ile)")
    p.add_argument("--kirilim", nargs="*", help="Çapraz tablo sütunları (varsayılan: türü Demografi olan sorular)")
    p.add_argument("--min-grup", type=int, default=30, help="Bu sayının altındaki kırılım grupları için uyarı (varsayılan 30)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "anket_analizi.xlsx")
    a = p.parse_args(argv)
    ornek = a.yanitlar == ORNEK / "yanitlar.csv"
    kod = a.kodlar or (ORNEK / "kod_cercevesi.csv" if ornek else None)
    for y in (a.yanitlar, a.sorular, kod):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        s = calistir(a.yanitlar, a.sorular, a.cikti, kod, a.kirilim, a.min_grup)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {s['n']} yanıt · {len(s['sorular'])} soru · genel hata payı ± %{s['hata_payi'] * 100:.1f}".replace(".", ","))
    for x in sorted(s["olcek"], key=lambda x: -x["ort"]):
        print(f"     {x['soru'].kod:<6} ort. {tr(x['ort'])} · üst 2 kutu %{x['ust2'] * 100:.0f} — {x['soru'].metin[:60]}")
    for x in sorted(s["anlamli"], key=lambda x: x["test"]["p"]):
        print(f"[i] Anlamlı fark: {x['soru'].kod} × {x['kirilim'].kod} (p = {p_yaz(x['test']['p'])}) — {x['ozet']}")
    for u in s["uyarilar"]:
        if u["onem"] != "Bilgi":
            print(f"[!] {u['kim']} · {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
