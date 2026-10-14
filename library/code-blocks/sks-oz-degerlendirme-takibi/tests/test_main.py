import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
GUN = date(2026, 9, 30)


def s(birim, kod, bolum, deger, puan, dok="", sorumlu="", termin=None):
    return {"kod": kod, "bolum": bolum, "birim": birim, "deger": deger, "puan": D(puan), "dokuman": "", "dok_durum": dok, "sorumlu": sorumlu, "termin": date(2026, *termin) if termin else None, "aciklama": ""}


def veri():
    return [s("B1", "S1", "X", "Karşılıyor", 2), s("B1", "S2", "X", "Kısmen", 2, sorumlu="Kalite"), s("B1", "S3", "Y", "Karşılamıyor", 1, "Yok", termin=(9, 1)),
            s("b1", "S4", "Y", "Değerlendirme dışı", 1), s("B2", "S1", "x", "Karşılıyor", 1, "Güncel değil", "A", (10, 15))]


def k(o):
    return (o["ad"], o["standart"], o["karsiliyor"], o["kismen"], o["karsilamiyor"], o["dd"], o["olasi"], o["kazanilan"], o["oran"], o["eksik_dok"])


class TakipTesti(unittest.TestCase):
    def test_puanlar(self):
        r = main.takip(veri(), GUN)
        self.assertEqual([k(b) for b in r["birim"]], [("B1", 4, 1, 1, 1, 1, D(5), D("3.00"), D("60.00"), 1), ("B2", 1, 1, 0, 0, 0, D(1), D("1.00"), D("100.00"), 1)])
        self.assertEqual([(b["ad"], b["oran"]) for b in r["bolum"]], [("X", D("80.00")), ("Y", D("0.00"))])
        self.assertEqual((r["toplam"]["oran"], [(a["birim"], a["kod"], a["neden"], a["gecikme"]) for a in r["aksiyon"]]),
                         (D("66.67"), [("B1", "S3", "Karşılamıyor, Doküman yok", 29), ("B2", "S1", "Doküman güncel değil", None), ("B1", "S2", "Kısmen", None)]))
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Bilgi", "Sorumlu yazılmamış", "B1 · S3"), ("Orta", "Termini geçmiş aksiyon", "B1 · S3"), ("Orta", "Uyum oranı hedefin altında", "B1")])

    def test_secenekler(self):
        r = main.takip(veri(), GUN, kismen=D(0), hedef=D(30))
        self.assertEqual(([b["oran"] for b in r["birim"]], [u["tur"] for u in r["uyarilar"] if u["onem"] == "Orta" and "Uyum" in u["tur"]]), ([D("40.00"), D("100.00")], []))
        self.assertEqual([main._deger(x) for x in ("Karşılıyor", "kısmen", "KARŞILAMIYOR", "Değerlendirme dışı", "evet", "hayır", "", "?")],
                         ["Karşılıyor", "Kısmen", "Karşılamıyor", "Değerlendirme dışı", "Karşılıyor", "Karşılamıyor", None, None])
        self.assertEqual([main._dok(x) for x in ("Var", "yok", "Güncel değil", "")], ["Var", "Yok", "Güncel değil", ""])
        with self.assertRaises(ValueError):
            main.takip([], GUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "oz_degerlendirme.csv", cikti, GUN)
            self.assertEqual([k(b) for b in r["birim"]] + [k({"ad": "T", **r["toplam"]})],
                             [("Acil", 12, 8, 1, 2, 1, D(125), D("95.00"), D("76.00"), 5), ("Genel Cerrahi", 12, 9, 1, 1, 1, D(95), D("80.00"), D("84.21"), 3),
                              ("Dahiliye", 12, 9, 1, 1, 1, D(105), D("92.50"), D("88.10"), 4), ("Yoğun Bakım", 12, 10, 0, 1, 1, D(110), D("100.00"), D("90.91"), 4),
                              ("Ameliyathane", 12, 12, 0, 0, 0, D(110), D("110.00"), D("100.00"), 5), ("T", 60, 48, 3, 5, 4, D(545), D("477.50"), D("87.61"), 21)])
            self.assertEqual((len(r["aksiyon"]), [u["tur"] for u in r["uyarilar"]].count("Termini geçmiş aksiyon"), [u["kim"] for u in r["uyarilar"] if u["tur"].startswith("Uyum")]), (25, 10, ["Acil"]))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Birimler", "Bölümler", "Aksiyon Listesi", "Tüm Değerlendirmeler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "d.csv"
            y.write_text("Standart No;Servis;Sonuç;Ağırlık;Hedef Tarih\nS1;A;Evet;3;\ns1;a;Hayır;;\nS2;A;belki;;\nS3;A;Evet;x;\nS4;A;Evet;;yok\nS5;;Kısmen;;01.10.2026\n;A;Evet;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["kod"], x["birim"], x["deger"], x["puan"], x["termin"]) for x in l], [u["tur"] for u in uy]),
                             ([("S1", "A", "Karşılıyor", D(3), None), ("S5", "Kurum geneli", "Kısmen", D(1), date(2026, 10, 1))],
                              ["Standart iki kez değerlendirilmiş", "Satır okunamadı", "Satır okunamadı", "Satır okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--degerlendirme", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "30.09.2026", "--kismen", "25", "--hedef", "90"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--kismen", "101"], ["--hedef", "-1"])], [2] * 3)
            self.assertEqual(main.main(["--degerlendirme", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Standart No;Servis\nS1;A\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
