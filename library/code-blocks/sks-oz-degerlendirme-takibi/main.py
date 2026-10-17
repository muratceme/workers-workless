"""
SKS Öz Değerlendirme Takibi — Gerisi Bende kod bloğu
Kalite Yönetimi › Kalite Birim Sorumlusu
SKS standartlarına göre öz değerlendirme puanlarını ve eksik dokümanları birim ve bölüm bazında takip eder.
  Kazanılan puan   = Karşılıyor: standart puanı · Kısmen: puan × --kismen (%50) · Karşılamıyor: 0 · Değerlendirme dışı: hesaba girmez
  Uyum oranı %     = Σ kazanılan puan ÷ Σ değerlendirilen standart puanı × 100   (birim, bölüm ve toplam)
  Eksik doküman    = Doküman Durumu "Yok" ya da "Güncel değil" olan satırlar
  Aksiyon listesi  = Kısmen / Karşılamıyor değerlendirilen standartlar ve eksik dokümanlar; sorumlu ve termin ile
  Uyarı            = birim uyum oranı --hedef (%80) altında · termini geçmiş aksiyon · sorumlusu yazılmamış aksiyon
Puanlama şeması kurumun öz değerlendirme yöntemine göre ayarlanmalıdır; kod resmî değerlendirme puanı üretmez.
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"kod": ("standart kodu", "standart no", "kod", "standart"), "bolum": ("bolum", "sks bolumu", "boyut", "baslik"), "birim": ("birim", "servis", "degerlendirilen birim"),
         "deger": ("degerlendirme", "durum", "sonuc", "uyum"), "puan": ("puan", "standart puani", "agirlik"), "dokuman": ("gerekli dokuman", "dokuman", "kanit"),
         "dok_durum": ("dokuman durumu", "kanit durumu", "dokuman var mi"), "sorumlu": ("sorumlu", "sorumlu kisi"), "termin": ("termin", "hedef tarih", "son tarih"),
         "aciklama": ("aciklama", "not", "bulgu")}
SIFIR = Decimal(0)


def _deger(x) -> str | None:
    k = katla(x)
    if not k:
        return None
    if k.startswith("kismen") or k in ("k", "kismi"):
        return "Kısmen"
    if k.startswith("karsilamiyor") or k in ("hayir", "h", "yok", "uygun degil", "karsilanmiyor"):
        return "Karşılamıyor"
    if k.startswith("karsiliyor") or k in ("evet", "e", "tam", "uygun", "karsilaniyor"):
        return "Karşılıyor"
    if "disi" in k or k in ("dd", "uygulanamaz", "na"):
        return "Değerlendirme dışı"
    return None


def _dok(x) -> str:
    k = katla(x)
    if not k:
        return ""
    if "guncel degil" in k or "eski" in k or "revize" in k:
        return "Güncel değil"
    if k in ("yok", "hayir", "eksik", "h"):
        return "Yok"
    return "Var"


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("kod", "deger")):
        kod = metin(r["kod"])
        if not kod:
            continue
        birim = metin(r["birim"]) or "Kurum geneli"
        d, puan = _deger(r["deger"]), para(r["puan"]) if metin(r["puan"]) else Decimal(1)
        termin = tarih(r["termin"]) if metin(r["termin"]) else None
        kim = f"{birim} · {kod}"
        if d is None or puan is None or puan < 0 or (metin(r["termin"]) and termin is None):
            uy.append(uyari("Yüksek", "Satır okunamadı", kim, f"Satır {r['_satir']}: değerlendirme ('{metin(r['deger'])}'), puan ya da termin okunamadı. Satır alınmadı."))
            continue
        if (katla(birim), katla(kod)) in gorulen:
            uy.append(uyari("Orta", "Standart iki kez değerlendirilmiş", kim, f"Satır {r['_satir']}: aynı birimde aynı standart ikinci kez geçiyor. İkinci satır alınmadı."))
            continue
        gorulen.add((katla(birim), katla(kod)))
        l.append({"kod": kod, "bolum": metin(r["bolum"]) or "Belirtilmemiş", "birim": birim, "deger": d, "puan": puan, "dokuman": metin(r["dokuman"]), "dok_durum": _dok(r["dok_durum"]),
                  "sorumlu": metin(r["sorumlu"]), "termin": termin, "aciklama": metin(r["aciklama"])})
    return l, uy


def _ozet(l: list[dict], kismen: Decimal) -> dict:
    deg = [x for x in l if x["deger"] != "Değerlendirme dışı"]
    olasi = sum((x["puan"] for x in deg), SIFIR)
    kaz = sum((x["puan"] if x["deger"] == "Karşılıyor" else x["puan"] * kismen / 100 if x["deger"] == "Kısmen" else SIFIR for x in deg), SIFIR)
    say = lambda d: sum(1 for x in l if x["deger"] == d)   # noqa: E731
    return {"standart": len(l), "karsiliyor": say("Karşılıyor"), "kismen": say("Kısmen"), "karsilamiyor": say("Karşılamıyor"), "dd": say("Değerlendirme dışı"),
            "olasi": olasi, "kazanilan": yuvarla(kaz), "oran": yuvarla(kaz * 100 / olasi) if olasi else None, "eksik_dok": sum(1 for x in l if x["dok_durum"] in ("Yok", "Güncel değil"))}


def takip(satirlar: list[dict], gun: date, kismen: Decimal = Decimal(50), hedef: Decimal = Decimal(80)) -> dict:
    if not satirlar:
        raise ValueError("Değerlendirme satırı yok.")
    uy, birimler, bolumler = [], {}, {}
    for x in satirlar:
        birimler.setdefault(katla(x["birim"]), []).append(x)
        bolumler.setdefault(katla(x["bolum"]), []).append(x)
    birim = sorted(({"ad": l[0]["birim"], **_ozet(l, kismen)} for l in birimler.values()), key=lambda b: (b["oran"] if b["oran"] is not None else Decimal(101), katla(b["ad"])))
    bolum = sorted(({"ad": l[0]["bolum"], **_ozet(l, kismen)} for l in bolumler.values()), key=lambda b: katla(b["ad"]))
    for b in birim:
        if b["oran"] is not None and b["oran"] < hedef:
            uy.append(uyari("Orta", "Uyum oranı hedefin altında", b["ad"], f"Uyum %{b['oran']} (hedef %{hedef}); {b['kismen']} kısmen, {b['karsilamiyor']} karşılamıyor."))
    aksiyon = []
    for x in satirlar:
        nedenler = [x["deger"]] if x["deger"] in ("Kısmen", "Karşılamıyor") else []
        if x["dok_durum"] in ("Yok", "Güncel değil"):
            nedenler.append(f"Doküman {x['dok_durum'].lower()}")
        if not nedenler:
            continue
        gecikme = (gun - x["termin"]).days if x["termin"] and x["termin"] < gun else None
        aksiyon.append({**x, "neden": ", ".join(nedenler), "gecikme": gecikme})
        kim = f"{x['birim']} · {x['kod']}"
        if gecikme:
            uy.append(uyari("Orta", "Termini geçmiş aksiyon", kim, f"Termin {x['termin']:%d.%m.%Y}, {gecikme} gün geçti ({', '.join(nedenler)})" + (f"; sorumlu {x['sorumlu']}." if x["sorumlu"] else ".")))
        if not x["sorumlu"]:
            uy.append(uyari("Bilgi", "Sorumlu yazılmamış", kim, ", ".join(nedenler) + "."))
    aksiyon.sort(key=lambda a: (a["termin"] or date.max, katla(a["birim"]), katla(a["kod"])))
    return {"toplam": _ozet(satirlar, kismen), "birim": birim, "bolum": bolum, "aksiyon": aksiyon, "satirlar": satirlar, "uyarilar": uy}


DEGER_RENK = {"Karşılıyor": YESIL, "Kısmen": SARI, "Karşılamıyor": KIRMIZI, "Değerlendirme dışı": None}


def rapor_yaz(cikti: Path, r: dict, gun: date, kismen: Decimal, hedef: Decimal) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Birimler"
    bas = ["Standart", "Karşılıyor", "Kısmen", "Karşılamıyor", "Değ. Dışı", "Olası Puan", "Kazanılan", "Uyum %", "Eksik Doküman"]
    _baslik(ws, ["Birim"] + bas, (24, 9, 10, 8, 11, 8, 10, 10, 8, 10), "B2")

    def yaz(ws, ad, o, kalin=False):
        k = _satir(ws, [ad, o["standart"], o["karsiliyor"], o["kismen"], o["karsilamiyor"], o["dd"], o["olasi"], o["kazanilan"], o["oran"], o["eksik_dok"]], para_s=(7, 8, 9),
                   dolgu={9: None if o["oran"] is None else KIRMIZI if o["oran"] < hedef else YESIL})
        if kalin:
            for h in ws[k]:
                h.font = Font(bold=True)
    for b in r["birim"]:
        yaz(ws, b["ad"], b)
    yaz(ws, "TOPLAM", r["toplam"], True)
    ws.append([])
    ws.append([f"Kısmen = puan × %{kismen}. Hedef %{hedef}. Rapor tarihi {gun:%d.%m.%Y}. Kurum içi öz değerlendirme içindir; resmî değerlendirme puanı değildir."])
    ws = wb.create_sheet("Bölümler")
    _baslik(ws, ["Bölüm"] + bas, (30, 9, 10, 8, 11, 8, 10, 10, 8, 10), "B2")
    for b in r["bolum"]:
        yaz(ws, b["ad"], b)
    ws = wb.create_sheet("Aksiyon Listesi")
    _baslik(ws, ["Termin", "Gecikme (gün)", "Birim", "Standart", "Bölüm", "Değerlendirme", "Neden", "Gerekli Doküman", "Sorumlu", "Açıklama"], (11, 9, 20, 12, 22, 13, 28, 26, 16, 30), "C2")
    for a in r["aksiyon"]:
        _satir(ws, [a["termin"], a["gecikme"], a["birim"], a["kod"], a["bolum"], a["deger"], a["neden"], a["dokuman"], a["sorumlu"], a["aciklama"]], tarih_s=(1,),
               dolgu={2: KIRMIZI if a["gecikme"] else None, 6: DEGER_RENK[a["deger"]]}, sar=(7, 10))
    ws = wb.create_sheet("Tüm Değerlendirmeler")
    _baslik(ws, ["Birim", "Bölüm", "Standart", "Değerlendirme", "Puan", "Doküman Durumu", "Sorumlu", "Termin"], (20, 22, 12, 13, 7, 12, 16, 11), "D2")
    for x in sorted(r["satirlar"], key=lambda x: (katla(x["birim"]), katla(x["kod"]))):
        _satir(ws, [x["birim"], x["bolum"], x["kod"], x["deger"], x["puan"], x["dok_durum"], x["sorumlu"], x["termin"]], tarih_s=(8,), dolgu={4: DEGER_RENK[x["deger"]]})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, gun: date, kismen: Decimal = Decimal(50), hedef: Decimal = Decimal(80)) -> dict:
    satirlar, uy = oku(yol)
    r = takip(satirlar, gun, kismen, hedef)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, gun, kismen, hedef)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="SKS öz değerlendirme puanlarını ve eksik dokümanları birim bazında takip eder.")
    p.add_argument("--degerlendirme", type=Path, help="Standart Kodu, Bölüm, Birim, Değerlendirme, Puan, Gerekli Doküman, Doküman Durumu, Sorumlu, Termin, Açıklama")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 30.09.2026)")
    p.add_argument("--kismen", default="50", help="Kısmen karşılanan standardın puan yüzdesi (varsayılan 50)")
    p.add_argument("--hedef", default="80", help="Birim uyum oranı bu yüzdenin altındaysa uyarılır (varsayılan 80)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "sks_oz_degerlendirme.xlsx")
    a = p.parse_args(argv)
    yol = a.degerlendirme or ORNEK / "oz_degerlendirme.csv"
    gun = tarih(a.tarih) if a.tarih else (date(2026, 9, 30) if a.degerlendirme is None else date.today())
    kismen, hedef = para(a.kismen), para(a.hedef)
    if gun is None or kismen is None or not 0 <= kismen <= 100 or hedef is None or not 0 <= hedef <= 100:
        print("[X] --tarih GG.AA.YYYY, --kismen ve --hedef 0-100 arası olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, gun, kismen, hedef)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {len(r['birim'])} birim · {t['standart']} değerlendirme · uyum %{t['oran']} · kısmen {t['kismen']} · karşılamıyor {t['karsilamiyor']} · eksik doküman {t['eksik_dok']} · aksiyon {len(r['aksiyon'])}")
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
