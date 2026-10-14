import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
V = [f"NM0AAAAAAAAAAAA{i:02d}" for i in range(10)]


def m(s, son=None, tel="0532"):
    return {"sasi": s, "plaka": f"P{s[-2:]}", "musteri": "M", "telefon": tel, "son": son}


class EslestirmeTesti(unittest.TestCase):
    def test_eslestir(self):
        k = {"K1": {"kod": "K1", "ad": "Röle", "bas": date(2026, 5, 1), "sasiler": V[:6]}, "K2": {"kod": "K2", "ad": "Yazılım", "bas": None, "sasiler": V[4:8]}}
        mus = {V[0]: m(V[0], date(2026, 6, 1)), V[1]: m(V[1], date(2026, 4, 1)), V[2]: m(V[2], date(2026, 7, 1), ""), V[4]: m(V[4]), V[6]: m(V[6], date(2026, 9, 1)), V[9]: m(V[9])}
        yap = {(V[0], "K1"): {}, (V[3], "K1"): {}, (V[6], "K2"): {}, (V[9], "K1"): {}}
        r = main.eslestir(k, mus, yap, date(2026, 10, 14))
        o = {s["kod"]: s for s in r["ozet"]}
        self.assertEqual([(o["K1"][a]) for a in ("toplam", "eslesen", "yapilan", "bizde_yapilan", "bekleyen", "kayitsiz", "tamamlanma", "bizde_tamamlanma")],
                         [6, 4, 2, 1, 3, 2, D("33.33"), D("25.00")])
        self.assertEqual([(o["K2"][a]) for a in ("toplam", "eslesen", "yapilan", "bekleyen", "tamamlanma")], [4, 2, 1, 1, D("25.00")])
        self.assertEqual([(b["kod"], b["sasi"][-2:], b["kacan"], b["bekleme"]) for b in r["bekleyen"]],
                         [("K1", "02", True, 166), ("K1", "01", False, 166), ("K1", "04", False, 166), ("K2", "04", False, None)])
        self.assertEqual(sorted((u["onem"], u["tur"]) for u in r["uyarilar"]),
                         [("Bilgi", "Listede olmayan işlem"), ("Bilgi", "Telefonu olmayan bekleyen araç"), ("Orta", "Kaçırılan ziyaret")])
        with self.assertRaises(ValueError):
            main.eslestir({}, mus, yap)

    def test_sasi(self):
        self.assertEqual(main.sasi(" nm0 abc-def "), "NM0ABCDEF")
        self.assertIsNone(main._bicim_uyari("NM0AAAAAAAAAAAA01", "x"))
        self.assertEqual([main._bicim_uyari(s, "x")["tur"] for s in ("NM0", "NM0AAAAAAAAAAAAO1", "NM0AAAAAAAAAAAA.1")], ["Şasi numarası biçimi"] * 3)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "kampanya_sasileri.csv", ORNEK / "musteriler.csv", ORNEK / "yapilanlar.csv", cikti, date(2026, 10, 14))
            self.assertEqual([(s["kod"], s["toplam"], s["eslesen"], s["yapilan"], s["bekleyen"]) for s in r["ozet"]],
                             [("GC-2026-07", 180, 112, 70, 42), ("GC-2026-11", 141, 83, 34, 49), ("SK-2026-03", 120, 73, 18, 55)])
            self.assertEqual(sum(1 for b in r["bekleyen"] if b["kacan"]), 35)
            self.assertEqual(sorted({u["tur"] for u in r["uyarilar"] if u["onem"] == "Bilgi"}),
                             ["Listede olmayan işlem", "Telefonu olmayan bekleyen araç", "Şasi kampanyada iki kez", "Şasi numarası biçimi"])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Kampanya Özeti", "Bekleyen Araçlar", "Uyarılar"])
            self.assertEqual(len((Path(tmp) / "t_bekleyen.csv").read_text(encoding="utf-8-sig").splitlines()), 147)

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            ky, my, yy = Path(tmp) / "k.csv", Path(tmp) / "m.csv", Path(tmp) / "y.csv"
            ky.write_text(f"Kampanya;Başlangıç;VIN\nk1;01.05.2026;{V[0]}\nK1;;{V[1]}\n;;{V[2]}\nK2;x;{V[3]}\n", encoding="utf-8")
            my.write_text(f"VIN;Plaka;Son Servis\n{V[0]};34 a 1;01.04.2026\n{V[0].lower()};34 A 1;01.06.2026\n{V[1]};;x\n", encoding="utf-8")
            yy.write_text(f"VIN;Kampanya;Tarih\n{V[1]};k1;01.06.2026\n;K1;\n{V[2]};K1;x\n", encoding="utf-8")
            k, uy = main.kampanya_oku(ky)
            self.assertEqual((k["K1"]["sasiler"], k["K1"]["bas"], len(uy)), (V[:2], date(2026, 5, 1), 2))
            mu, uy = main.musteri_oku(my)
            self.assertEqual((mu[V[0]]["son"], mu[V[0]]["plaka"], len(uy)), (date(2026, 6, 1), "34 A 1", 1))
            y, uy = main.yapilan_oku(yy)
            self.assertEqual((list(y), len(uy)), ([(V[1], "K1")], 2))
            c = str(Path(tmp) / "x.xlsx")
            self.assertEqual(main.main(["--kampanya", str(ky), "--musteriler", str(my), "--yapilanlar", str(yy), "--tarih", "14.10.2026", "--cikti", c]), 0)
            self.assertEqual(main.main(["--kampanya", str(ky), "--musteriler", str(my), "--cikti", c]), 0)
            self.assertEqual(main.main(["--tarih", "x", "--cikti", c]), 2)
            self.assertEqual(main.main(["--kampanya", str(Path(tmp) / "yok.csv"), "--musteriler", str(my)]), 1)
            ky.write_text("Kampanya;VIN\n;\n", encoding="utf-8")
            self.assertEqual(main.main(["--kampanya", str(ky), "--musteriler", str(my), "--cikti", c]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
