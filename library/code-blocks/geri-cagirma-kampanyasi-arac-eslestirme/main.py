"""
Geri Çağırma Kampanyası Araç Eşleştirme — Workers / Workless kod bloğu
Satış Sonrası Hizmetler (Servis) › Garanti Uzmanı
Markanın kampanya şasi listesini servis müşteri kayıtlarıyla eşleştirip kampanyası yapılmamış araçları listeler.
  Şasi normalizasyonu = büyük harf, boşluk ve tire atılır
  Eşleşen             = kampanya şasisi servis kayıtlarında var
  Yapılan             = şasi + kampanya kodu yapılan işlemler listesinde var
  Bekleyen            = eşleşen ve yapılmamış; son servis tarihi kampanya başlangıcından sonraysa "kaçırılan ziyaret"
  Tamamlanma %        = yapılan ÷ kampanya şasi sayısı × 100;  bizdeki tamamlanma % = yapılan ÷ eşleşen × 100
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
KAMPANYA = {"kod": ("kampanya kodu", "kampanya no", "kampanya"), "ad": ("kampanya adi", "aciklama", "konu"), "bas": ("baslangic tarihi", "baslangic", "yayin tarihi"),
            "sasi": ("sasi no", "sasi", "vin", "sase no")}
MUSTERI = {"sasi": ("sasi no", "sasi", "vin", "sase no"), "plaka": ("plaka",), "musteri": ("musteri", "ad soyad", "musteri adi"), "telefon": ("telefon", "gsm", "cep telefonu"),
           "son": ("son servis tarihi", "son servis", "son giris tarihi", "son ziyaret")}
YAPILAN = {"sasi": ("sasi no", "sasi", "vin", "sase no"), "kod": ("kampanya kodu", "kampanya no", "kampanya"), "tarih": ("islem tarihi", "tarih", "yapilis tarihi"),
           "is_emri": ("is emri no", "is emri")}


def sasi(x) -> str:
    return re.sub(r"[\s\-]+", "", metin(x).upper())


def _bicim_uyari(s: str, kim: str) -> dict | None:
    if len(s) != 17 or any(c in "IOQ" for c in s) or not s.isalnum():
        return uyari("Bilgi", "Şasi numarası biçimi", kim, f"'{s}' 17 karakter olmalı, yalnız harf / rakam içermeli ve I, O, Q harflerini içermemeli.")
    return None


def kampanya_oku(yol: Path) -> tuple[dict, list[dict]]:
    uy, k = [], {}
    for r in kayitlar(yol, KAMPANYA, ("kod", "sasi")):
        kod, s = metin(r["kod"]), sasi(r["sasi"])
        if not kod and not s:
            continue
        b = tarih(r["bas"]) if metin(r["bas"]) else None
        if not kod or not s or (metin(r["bas"]) and b is None):
            uy.append(uyari("Yüksek", "Kampanya satırı okunamadı", kod or s, f"Satır {r['_satir']}: kampanya kodu, şasi ya da başlangıç tarihi okunamadı."))
            continue
        c = k.setdefault(kod.upper(), {"kod": kod.upper(), "ad": metin(r["ad"]), "bas": b, "sasiler": []})
        c["ad"], c["bas"] = c["ad"] or metin(r["ad"]), c["bas"] or b
        if s in c["sasiler"]:
            uy.append(uyari("Bilgi", "Şasi kampanyada iki kez", f"{kod} · {s}", f"Satır {r['_satir']}: tekrar alınmadı."))
            continue
        b_uy = _bicim_uyari(s, f"{kod} · {s}")
        if b_uy:
            uy.append(b_uy)
        c["sasiler"].append(s)
    return k, uy


def musteri_oku(yol: Path) -> tuple[dict, list[dict]]:
    uy, d = [], {}
    for r in kayitlar(yol, MUSTERI, ("sasi",)):
        s = sasi(r["sasi"])
        if not s:
            continue
        son = tarih(r["son"]) if metin(r["son"]) else None
        if metin(r["son"]) and son is None:
            uy.append(uyari("Orta", "Son servis tarihi okunamadı", s, f"Satır {r['_satir']}: tarih boş sayıldı."))
        kayit = {"sasi": s, "plaka": " ".join(metin(r["plaka"]).upper().split()), "musteri": metin(r["musteri"]), "telefon": metin(r["telefon"]), "son": son}
        eski = d.get(s)
        if eski is None or (kayit["son"] and (eski["son"] is None or kayit["son"] > eski["son"])):
            d[s] = kayit
    return d, uy


def yapilan_oku(yol: Path) -> tuple[dict, list[dict]]:
    uy, d = [], {}
    for r in kayitlar(yol, YAPILAN, ("sasi", "kod")):
        s, kod = sasi(r["sasi"]), metin(r["kod"]).upper()
        if not s and not kod:
            continue
        t = tarih(r["tarih"]) if metin(r["tarih"]) else None
        if not s or not kod or (metin(r["tarih"]) and t is None):
            uy.append(uyari("Yüksek", "İşlem satırı okunamadı", s or kod, f"Satır {r['_satir']}: şasi, kampanya kodu ya da işlem tarihi okunamadı."))
            continue
        d.setdefault((s, kod), {"tarih": t, "is_emri": metin(r["is_emri"])})
    return d, uy


def eslestir(kampanyalar: dict, musteriler: dict, yapilanlar: dict, bugun: date | None = None) -> dict:
    if not kampanyalar:
        raise ValueError("Okunabilen kampanya şasisi yok.")
    uy, ozet, bekleyen = [], [], []
    for c in sorted(kampanyalar.values(), key=lambda c: c["kod"]):
        eslesen = [s for s in c["sasiler"] if s in musteriler]
        yapildi = [s for s in c["sasiler"] if (s, c["kod"]) in yapilanlar]
        bizde_yapilan = [s for s in eslesen if (s, c["kod"]) in yapilanlar]
        for s in eslesen:
            if (s, c["kod"]) in yapilanlar:
                continue
            m = musteriler[s]
            kacan = bool(c["bas"] and m["son"] and m["son"] >= c["bas"])
            bekleyen.append({"kod": c["kod"], "ad": c["ad"], "bas": c["bas"], **m, "kacan": kacan, "bekleme": (bugun - c["bas"]).days if bugun and c["bas"] else None})
            if kacan:
                uy.append(uyari("Orta", "Kaçırılan ziyaret", f"{c['kod']} · {m['plaka'] or s}", f"Araç kampanya başladıktan sonra ({c['bas']:%d.%m.%Y}) {m['son']:%d.%m.%Y} tarihinde servise geldi ama kampanya yapılmadı."))
        ozet.append({"kod": c["kod"], "ad": c["ad"], "bas": c["bas"], "toplam": len(c["sasiler"]), "eslesen": len(eslesen), "yapilan": len(yapildi), "bizde_yapilan": len(bizde_yapilan),
                     "bekleyen": len(eslesen) - len(bizde_yapilan), "kayitsiz": len(c["sasiler"]) - len(eslesen),
                     "tamamlanma": yuvarla(Decimal(len(yapildi)) * 100 / len(c["sasiler"])) if c["sasiler"] else None,
                     "bizde_tamamlanma": yuvarla(Decimal(len(bizde_yapilan)) * 100 / len(eslesen)) if eslesen else None})
    kampanya_sasi = {(s, c["kod"]) for c in kampanyalar.values() for s in c["sasiler"]}
    disari = sorted(k for k in yapilanlar if k not in kampanya_sasi)
    for s, kod in disari:
        uy.append(uyari("Bilgi", "Listede olmayan işlem", f"{kod} · {s}", "Kampanya işlemi yapılmış görünüyor ama bu şasi kampanya listesinde yok; kod ya da şasi yazımını kontrol edin."))
    telsiz = sum(1 for b in bekleyen if not b["telefon"])
    if telsiz:
        uy.append(uyari("Bilgi", "Telefonu olmayan bekleyen araç", f"{telsiz} araç", "Bekleyen araçların bir kısmında telefon yok; adres / e-posta ya da marka çağrı merkezi üzerinden ulaşın."))
    bekleyen.sort(key=lambda b: (not b["kacan"], b["kod"], b["son"] is None, -(b["son"].toordinal()) if b["son"] else 0, b["sasi"]))
    return {"ozet": ozet, "bekleyen": bekleyen, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kampanya Özeti"
    _baslik(ws, ["Kampanya", "Açıklama", "Başlangıç", "Kampanya Şasisi", "Bizde Kayıtlı", "Yapılan (toplam)", "Bizde Yapılan", "Bizde Bekleyen", "Kaydımızda Yok", "Tamamlanma %",
                 "Bizdeki Tamamlanma %"], (12, 32, 11, 10, 10, 10, 10, 10, 10, 11, 12))
    for s in r["ozet"]:
        _satir(ws, [s["kod"], s["ad"], s["bas"], s["toplam"], s["eslesen"], s["yapilan"], s["bizde_yapilan"], s["bekleyen"], s["kayitsiz"], s["tamamlanma"], s["bizde_tamamlanma"]],
               para_s=(10, 11), tarih_s=(3,), sar=(2,))
    ws.append([])
    ws.append(["Kaydımızda yok: kampanya listesinde olup servis müşteri kayıtlarında bulunmayan şasiler (başka yetkili serviste işlem görüyor olabilir)."])
    ws = wb.create_sheet("Bekleyen Araçlar")
    _baslik(ws, ["Kampanya", "Şasi No", "Plaka", "Müşteri", "Telefon", "Son Servis", "Kaçırılan Ziyaret", "Kampanya Başlangıcından Beri (gün)"], (12, 20, 12, 22, 15, 11, 10, 12), "C2")
    for b in r["bekleyen"]:
        _satir(ws, [b["kod"], b["sasi"], b["plaka"], b["musteri"], b["telefon"], b["son"], "Evet" if b["kacan"] else "", b["bekleme"]], tarih_s=(6,), dolgu={7: SARI if b["kacan"] else None})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(kampanya_yol: Path, musteri_yol: Path, yapilan_yol: Path | None, cikti: Path, bugun: date | None = None) -> dict:
    k, uy1 = kampanya_oku(kampanya_yol)
    m, uy2 = musteri_oku(musteri_yol)
    y, uy3 = yapilan_oku(yapilan_yol) if yapilan_yol else ({}, [])
    r = eslestir(k, m, y, bugun)
    r["uyarilar"] = sirala(uy1 + uy2 + uy3 + r["uyarilar"])
    rapor_yaz(cikti, r)
    with open(cikti.with_name(cikti.stem + "_bekleyen.csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Kampanya", "Şasi No", "Plaka", "Müşteri", "Telefon", "Son Servis", "Kaçırılan Ziyaret"])
        for b in r["bekleyen"]:
            w.writerow([b["kod"], b["sasi"], b["plaka"], b["musteri"], b["telefon"], f"{b['son']:%d.%m.%Y}" if b["son"] else "", "Evet" if b["kacan"] else ""])
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Kampanya şasi listesini servis kayıtlarıyla eşleştirip kampanyası yapılmamış araçları listeler.")
    p.add_argument("--kampanya", type=Path, help="Kampanya Kodu, Kampanya Adı, Başlangıç Tarihi, Şasi No")
    p.add_argument("--musteriler", type=Path, help="Şasi No, Plaka, Müşteri, Telefon, Son Servis Tarihi")
    p.add_argument("--yapilanlar", type=Path, help="Şasi No, Kampanya Kodu, İşlem Tarihi, İş Emri No (verilmezse hiçbiri yapılmamış sayılır)")
    p.add_argument("--tarih", help="Rapor günü GG.AA.YYYY (varsayılan bugün; örnek veride 14.10.2026)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "geri_cagirma_eslestirme.xlsx")
    a = p.parse_args(argv)
    ornek = a.kampanya is None and a.musteriler is None
    kamp, mus = a.kampanya or ORNEK / "kampanya_sasileri.csv", a.musteriler or ORNEK / "musteriler.csv"
    yap = a.yapilanlar or (ORNEK / "yapilanlar.csv" if ornek else None)
    bugun = tarih(a.tarih) if a.tarih else (date(2026, 10, 14) if ornek else date.today())
    if bugun is None:
        print("[X] --tarih GG.AA.YYYY olmalı")
        return 2
    for y in (kamp, mus, yap):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(kamp, mus, yap, a.cikti, bugun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    for s in r["ozet"]:
        print(f"[OK] {s['kod']}: {s['toplam']} şasi · bizde {s['eslesen']} · yapılan {s['yapilan']} (%{s['tamamlanma']}) · bizde bekleyen {s['bekleyen']}")
    print(f"     Kaçırılan ziyaret: {sum(1 for b in r['bekleyen'] if b['kacan'])} araç")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
