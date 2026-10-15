import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

O = main.ORNEK
YOLLAR = (O / "fatura_kalemleri.csv", O / "anlasma_fiyatlari.csv", O / "provizyonlar.csv", O / "policeler.csv")


def kalem(**kw):
    k = {"satir": 2, "dosya": "D1", "police": "P1", "provizyon": "R1", "tarih": date(2026, 10, 5), "kod": "A", "ad": "x", "adet": D(1), "fiyat": D(100), "tutar": D(100)}
    return k | kw


AN = {"A": {"ad": "A", "fiyat": D(100), "provizyon": True, "sinir": D(2)}}
PR = {"R1": {"dosya": "D1", "tarih": date(2026, 10, 1), "gun": 10, "tutar": D(250), "kodlar": {"A"}}}
PO = {"P1": {"bas": date(2026, 1, 1), "bit": date(2026, 12, 31), "limit": D(1000), "kullanilan": D(900), "katilim": D("0.2"), "disi": set()}}


class FaturaTesti(unittest.TestCase):
    def kontrol(self, kal, an=AN, pr=PR, po=PO):
        main.kalem_kontrol(kal, an, pr, po)
        return kal

    def test_yardimcilar(self):
        self.assertEqual((main.para("1.250,50"), main.para("750.000"), main.para("1250.5"), main.para("")), (D("1250.50"), D("750000"), D("1250.5"), None))
        self.assertEqual(main.tl(D("4200")), "4.200,00")
        self.assertEqual(main.kodlar("H-100 H-210;h-300"), {"H-100", "H-210", "H-300"})

    def test_kalem_kurallari(self):
        self.assertEqual(self.kontrol([kalem()])[0]["nedenler"], [])
        k = self.kontrol([kalem(fiyat=D(120), tutar=D(120))])[0]
        self.assertEqual((k["risk"], k["nedenler"]), (D("20.00"), ["Fiyat farkı: birim 120,00, anlaşma 100,00"]))
        self.assertEqual(self.kontrol([kalem(provizyon="")])[0]["risk"], D("100.00"))
        self.assertIn("geçerliliği dışında", self.kontrol([kalem(tarih=date(2026, 10, 12))])[0]["nedenler"][0])
        self.assertIn("onaylı değil", self.kontrol([kalem()], pr={"R1": PR["R1"] | {"kodlar": {"B"}}})[0]["nedenler"][0])
        self.assertIn("başka dosyaya", self.kontrol([kalem(dosya="D9")])[0]["nedenler"][0])
        self.assertIn("poliçe süresi", self.kontrol([kalem(tarih=date(2026, 10, 5))], po={"P1": PO["P1"] | {"bit": date(2026, 9, 30)}})[0]["nedenler"][0])
        self.assertIn("kapsam dışı", self.kontrol([kalem()], po={"P1": PO["P1"] | {"disi": {"A"}}})[0]["nedenler"][0])
        self.assertEqual(self.kontrol([kalem(kod="Z")])[0]["nedenler"], ["Anlaşma dışı hizmet kodu"])
        k = self.kontrol([kalem(adet=D(3), tutar=D(300))])[0]
        self.assertEqual(k["risk"], D("100.00"))  # 1 adet sınır aşımı
        k = self.kontrol([kalem(adet=D(2), tutar=D(150))])[0]
        self.assertEqual((k["risk"], k["nedenler"][0][:6]), (D("50.00"), "Hesap:"))
        k = self.kontrol([kalem(), kalem()])
        self.assertEqual((k[0]["risk"], k[1]["nedenler"]), (D("0.00"), ["Mükerrer kalem (aynı dosya, kod, tarih ve fiyat)"]))

    def test_dosya_kurallari(self):
        kal = self.kontrol([kalem(adet=D(2), tutar=D(200)), kalem(kod="Z", tutar=D(100))])
        d = main.dosya_kontrol(kal, PR, PO)[0]
        self.assertEqual((d["toplam"], d["sirket"], d["katilim"]), (D(300), D("240.00"), D("60.00")))
        self.assertEqual(d["notlar"], ["Dosya toplamı (300,00) provizyon onaylı tutarı (250,00) aşıyor", "Şirket payı (240,00) kalan limiti (100,00) aşıyor; aşan 140,00"])

    def test_ornek_ve_rapor(self):
        with tempfile.TemporaryDirectory() as tmp:
            r = main.calistir(*YOLLAR, Path(tmp) / "o.xlsx")
            risk = {(k["dosya"], k["kod"], i): k["risk"] for i, k in enumerate(r["kalemler"])}
            self.assertEqual(sum(risk.values()), D("21630.00"))
            self.assertEqual([k["risk"] for k in r["kalemler"]], [D(0), D(0), D(400), D(0), D(3800), D(7800), D(0), D(180), D(9000), D(450)])
            self.assertEqual([len(d["notlar"]) for d in r["dosyalar"]], [0, 1, 1, 1])
            ws = load_workbook(Path(tmp) / "o.xlsx")["Kalemler"]
            self.assertEqual((ws["E7"].value, ws["J7"].value), ("H-400", 7800))
            self.assertIn("Adet sınırı 10 aşıldı", ws["K7"].value)
            self.assertEqual(load_workbook(Path(tmp) / "o.xlsx")["Özet"]["B2"].value, 21630)

    def test_cli_ve_hatalar(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "c.xlsx")]), 0)
            self.assertEqual(main.main(["--police", str(Path(tmp) / "yok.csv")]), 1)
            bozuk = Path(tmp) / "f.csv"
            bozuk.write_text("Dosya No;Hizmet Tarihi;Hizmet Kodu;Adet;Birim Fiyat\nD1;dün;A;1;100\n", encoding="utf-8")
            self.assertEqual(main.main(["--fatura", str(bozuk)]), 1)


if __name__ == "__main__":
    unittest.main()
