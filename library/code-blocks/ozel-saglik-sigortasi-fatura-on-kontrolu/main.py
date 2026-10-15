"""
Özel Sağlık Sigortası Fatura Ön Kontrolü — Workers / Workless kod bloğu
Anlaşmalı Kurumlar ve Faturalama › Anlaşmalı Kurumlar Uzmanı
Fatura kalemlerini sigorta şirketi anlaşma fiyatları, provizyon onayı ve poliçe bilgileriyle karşılaştırıp kesinti riskli
kalemleri işaretler.

  Kalem kontrolleri:
    Hesap         = adet × birim fiyat ≠ tutar (±0,01)
    Anlaşma dışı  = hizmet kodu anlaşma listesinde yok → tutarın tamamı risk
    Fiyat farkı   = birim fiyat > anlaşma fiyatı → (birim − anlaşma) × adet risk
    Provizyon     = provizyon gerekli kalemde provizyon yok / dosyaya ait değil / hizmet tarihi onay + geçerlilik günü dışında /
                    kod onaylı kodlarda yok → tutarın tamamı risk
    Poliçe        = hizmet tarihi poliçe başlangıç-bitiş dışında ya da kod poliçede kapsam dışı → tutarın tamamı risk
    Adet sınırı   = dosyada aynı kodun toplam adedi > sınır → aşan adet × birim fiyat risk
    Mükerrer      = aynı dosya, kod, tarih, birim fiyatta ikinci ve sonraki satırlar → tutar risk
  Dosya kontrolleri:
    Provizyon tutarı = dosya toplamı > onaylı tutar
    Limit            = şirket payı (toplam × (1 − katılım payı)) > kalan limit (yıllık limit − kullanılan)
  Bir kalemde birden fazla neden varsa risk tutarı en yüksek nedene göre alınır (kalem tutarını aşmaz).
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
KIRMIZI, SARI, YESIL = "FDE2E1", "FFF4CE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")
SIFIR = Decimal(0)

KALEM = {"dosya": ("Dosya No",), "police": ("Poliçe No",), "provizyon": ("Provizyon No",), "tarih": ("Hizmet Tarihi", "Tarih"),
         "kod": ("Hizmet Kodu", "Kod"), "ad": ("Hizmet Adı",), "adet": ("Adet", "Miktar"), "fiyat": ("Birim Fiyat",), "tutar": ("Tutar",)}
ANLASMA = {"kod": ("Hizmet Kodu", "Kod"), "ad": ("Hizmet Adı",), "fiyat": ("Anlaşma Fiyatı", "Fiyat"), "provizyon": ("Provizyon Gerekli",),
           "sinir": ("Adet Sınırı",)}
PROVIZYON = {"no": ("Provizyon No",), "dosya": ("Dosya No",), "tarih": ("Onay Tarihi",), "gun": ("Geçerlilik Günü", "Geçerlilik"),
             "tutar": ("Onaylı Tutar",), "kodlar": ("Onaylı Hizmet Kodları", "Onaylı Kodlar")}
POLICE = {"no": ("Poliçe No",), "bas": ("Başlangıç",), "bit": ("Bitiş",), "limit": ("Yıllık Limit", "Limit"), "kullanilan": ("Kullanılan Limit",),
          "katilim": ("Katılım Payı %", "Katılım Payı"), "disi": ("Kapsam Dışı Kodlar", "Muafiyet Kodları")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").replace("İ", "i").replace("I", "ı").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 yazımlarını Decimal'e çevirir."""
    if x is None or str(x).strip() == "":
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("%", "").replace(" ", "")
    if re.fullmatch(r"-?\d{1,3}(\.\d{3})+(,\d+)?", s) or "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tl(x) -> str:
    return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def evet(x) -> bool:
    return katla(x) in {"evet", "e", "var", "1", "x"}


def kodlar(x) -> set[str]:
    return {k.strip().upper() for k in re.split(r"[\s;,|]+", str(x or "")) if k.strip()}


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    if not s:
        raise ValueError(f"{yol.name}: tablo boş")
    b = [katla(x) for x in s[0]]
    idx = {a: next((i for i, x in enumerate(b) if x in {katla(e) for e in es}), None) for a, es in sozluk.items()}
    eksik = [sozluk[a][0] for a in zorunlu if idx[a] is None]
    if eksik:
        raise ValueError(f"{yol.name}: gerekli sütun yok: {', '.join(eksik)}")
    return [{a: (r[i] if i is not None and i < len(r) and r[i] is not None else "") for a, i in idx.items()} for r in s[1:]]


def oku(fatura: Path, anlasma: Path, provizyon: Path, police: Path) -> tuple[list, dict, dict, dict]:
    kal = []
    for i, r in enumerate(kayitlar(fatura, KALEM, ("dosya", "tarih", "kod", "adet", "fiyat")), 2):
        t, a, f = tarih(r["tarih"]), para(r["adet"]), para(r["fiyat"])
        if not t or a is None or f is None:
            raise ValueError(f"{fatura.name} satır {i}: tarih / adet / birim fiyat okunamadı")
        tutar = para(r["tutar"])
        kal.append({"satir": i, "dosya": str(r["dosya"]).strip(), "police": str(r["police"]).strip(), "provizyon": str(r["provizyon"]).strip(),
                    "tarih": t, "kod": str(r["kod"]).strip().upper(), "ad": str(r["ad"]), "adet": a, "fiyat": f, "tutar": tutar if tutar is not None else a * f})
    an = {}
    for r in kayitlar(anlasma, ANLASMA, ("kod", "fiyat")):
        if str(r["kod"]).strip():
            an[str(r["kod"]).strip().upper()] = {"ad": str(r["ad"]), "fiyat": para(r["fiyat"]), "provizyon": evet(r["provizyon"]), "sinir": para(r["sinir"])}
    pr = {}
    for r in kayitlar(provizyon, PROVIZYON, ("no", "dosya", "tarih")):
        if str(r["no"]).strip():
            gun = para(r["gun"])
            pr[str(r["no"]).strip()] = {"dosya": str(r["dosya"]).strip(), "tarih": tarih(r["tarih"]), "gun": int(gun) if gun is not None else None,
                                         "tutar": para(r["tutar"]), "kodlar": kodlar(r["kodlar"])}
    po = {}
    for r in kayitlar(police, POLICE, ("no", "bas", "bit")):
        if str(r["no"]).strip():
            po[str(r["no"]).strip()] = {"bas": tarih(r["bas"]), "bit": tarih(r["bit"]), "limit": para(r["limit"]), "kullanilan": para(r["kullanilan"]) or SIFIR,
                                         "katilim": (para(r["katilim"]) or SIFIR) / 100, "disi": kodlar(r["disi"])}
    return kal, an, pr, po


def kalem_kontrol(kal: list, an: dict, pr: dict, po: dict) -> None:
    gorulen = set()
    adet_top = defaultdict(lambda: SIFIR)
    for k in kal:
        nedenler = []  # (açıklama, risk tutarı)
        if abs(k["adet"] * k["fiyat"] - k["tutar"]) > KURUS:
            nedenler.append((f"Hesap: {k['adet']} × {tl(k['fiyat'])} = {tl(k['adet'] * k['fiyat'])}, faturada {tl(k['tutar'])}", abs(k["adet"] * k["fiyat"] - k["tutar"])))
        a = an.get(k["kod"])
        if not a:
            nedenler.append(("Anlaşma dışı hizmet kodu", k["tutar"]))
        else:
            if a["fiyat"] is not None and k["fiyat"] > a["fiyat"] + KURUS:
                nedenler.append((f"Fiyat farkı: birim {tl(k['fiyat'])}, anlaşma {tl(a['fiyat'])}", (k["fiyat"] - a["fiyat"]) * k["adet"]))
            if a["provizyon"]:
                p = pr.get(k["provizyon"])
                if not k["provizyon"]:
                    nedenler.append(("Provizyon gerekli ama provizyon numarası yok", k["tutar"]))
                elif not p:
                    nedenler.append((f"Provizyon {k['provizyon']} kayıtlarda yok", k["tutar"]))
                else:
                    if p["dosya"] and p["dosya"] != k["dosya"]:
                        nedenler.append((f"Provizyon {k['provizyon']} başka dosyaya ait ({p['dosya']})", k["tutar"]))
                    if p["tarih"] and p["gun"] is not None and not p["tarih"] <= k["tarih"] <= p["tarih"] + timedelta(days=p["gun"]):
                        son = p["tarih"] + timedelta(days=p["gun"])
                        nedenler.append((f"Hizmet tarihi provizyon geçerliliği dışında ({p['tarih']:%d.%m.%Y}-{son:%d.%m.%Y})", k["tutar"]))
                    if p["kodlar"] and k["kod"] not in p["kodlar"]:
                        nedenler.append(("Hizmet kodu provizyonda onaylı değil", k["tutar"]))
            adet_top[(k["dosya"], k["kod"])] += k["adet"]
            if a["sinir"] is not None and adet_top[(k["dosya"], k["kod"])] > a["sinir"]:
                asan = min(k["adet"], adet_top[(k["dosya"], k["kod"])] - a["sinir"])
                nedenler.append((f"Adet sınırı {a['sinir']:g} aşıldı (dosyada toplam {adet_top[(k['dosya'], k['kod'])]:g})", asan * k["fiyat"]))
        p = po.get(k["police"])
        if not p:
            nedenler.append(("Poliçe kayıtlarda yok", k["tutar"]))
        else:
            if p["bas"] and p["bit"] and not p["bas"] <= k["tarih"] <= p["bit"]:
                nedenler.append((f"Hizmet tarihi poliçe süresi dışında ({p['bas']:%d.%m.%Y}-{p['bit']:%d.%m.%Y})", k["tutar"]))
            if k["kod"] in p["disi"]:
                nedenler.append(("Hizmet poliçede kapsam dışı", k["tutar"]))
        anahtar = (k["dosya"], k["kod"], k["tarih"], k["fiyat"])
        if anahtar in gorulen:
            nedenler.append(("Mükerrer kalem (aynı dosya, kod, tarih ve fiyat)", k["tutar"]))
        gorulen.add(anahtar)
        k["nedenler"] = [n for n, _ in nedenler]
        k["risk"] = min(max((r for _, r in nedenler), default=SIFIR), max(k["tutar"], SIFIR)).quantize(KURUS)


def dosya_kontrol(kal: list, pr: dict, po: dict) -> list[dict]:
    dosyalar = defaultdict(list)
    for k in kal:
        dosyalar[k["dosya"]].append(k)
    out = []
    for no, ks in dosyalar.items():
        top = sum((k["tutar"] for k in ks), SIFIR)
        risk = sum((k["risk"] for k in ks), SIFIR)
        notlar = []
        pno = next((k["provizyon"] for k in ks if k["provizyon"]), "")
        p = pr.get(pno)
        if p and p["tutar"] is not None and top > p["tutar"] + KURUS:
            notlar.append(f"Dosya toplamı ({tl(top)}) provizyon onaylı tutarı ({tl(p['tutar'])}) aşıyor")
        pol = po.get(ks[0]["police"])
        sirket = top
        if pol:
            sirket = (top * (1 - pol["katilim"])).quantize(KURUS)
            if pol["limit"] is not None:
                kalan = pol["limit"] - pol["kullanilan"]
                if sirket > kalan + KURUS:
                    notlar.append(f"Şirket payı ({tl(sirket)}) kalan limiti ({tl(kalan)}) aşıyor; aşan {tl(sirket - kalan)}")
        out.append({"dosya": no, "police": ks[0]["police"], "provizyon": pno, "kalem": len(ks), "toplam": top, "katilim": top - sirket, "sirket": sirket,
                    "risk": risk, "notlar": notlar})
    return out


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[1]:
        h.fill, h.font, h.alignment = BASLIK_DOLGU, BASLIK_YAZI, UST
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def rapor_yaz(cikti: Path, kal: list, dos: list) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kalemler"
    _baslik(ws, ["Dosya No", "Poliçe", "Provizyon", "Tarih", "Kod", "Hizmet", "Adet", "Birim Fiyat", "Tutar", "Riskli Tutar", "Nedenler"],
            (10, 10, 11, 11, 8, 28, 6, 12, 12, 12, 70))
    for k in kal:
        ws.append([k["dosya"], k["police"], k["provizyon"], k["tarih"], k["kod"], k["ad"], float(k["adet"]), float(k["fiyat"]), float(k["tutar"]),
                   float(k["risk"]), "\n".join(k["nedenler"])])
        ws.cell(ws.max_row, 4).number_format = TF
        for c in (8, 9, 10):
            ws.cell(ws.max_row, c).number_format = PF
        ws.cell(ws.max_row, 10).fill = PatternFill("solid", fgColor=KIRMIZI if k["risk"] > 0 else (SARI if k["nedenler"] else YESIL))
        for h in ws[ws.max_row]:
            h.alignment = UST
    d = wb.create_sheet("Dosyalar")
    _baslik(d, ["Dosya No", "Poliçe", "Provizyon", "Kalem", "Toplam", "Katılım Payı", "Şirket Payı", "Riskli Tutar", "Dosya Uyarıları"],
            (10, 10, 11, 7, 12, 12, 12, 12, 70))
    for x in dos:
        d.append([x["dosya"], x["police"], x["provizyon"], x["kalem"], float(x["toplam"]), float(x["katilim"]), float(x["sirket"]), float(x["risk"]),
                  "\n".join(x["notlar"])])
        for c in (5, 6, 7, 8):
            d.cell(d.max_row, c).number_format = PF
        if x["notlar"] or x["risk"] > 0:
            d.cell(d.max_row, 9).fill = PatternFill("solid", fgColor=KIRMIZI)
        for h in d[d.max_row]:
            h.alignment = UST
    o = wb.create_sheet("Özet")
    top = sum((k["tutar"] for k in kal), SIFIR)
    risk = sum((k["risk"] for k in kal), SIFIR)
    sayac = defaultdict(int)
    for k in kal:
        for n in k["nedenler"]:
            sayac[n.split(":")[0].split(" (")[0]] += 1
    o.append(["Toplam fatura", float(top)])
    o.append(["Kesinti riskli tutar", float(risk)])
    o.append(["Risk oranı %", float((risk / top * 100).quantize(KURUS)) if top else 0])
    o.append([])
    o.append(["Neden", "Kalem sayısı"])
    for n, s in sorted(sayac.items(), key=lambda x: -x[1]):
        o.append([n, s])
    o.append([])
    o.append(["Ön kontroldür. Kesin değerlendirme sigorta şirketinin provizyon ve fatura inceleme sonucuna göre yapılır; anlaşma fiyat listesinin "
              "güncel olduğunu kontrol edin."])
    o.column_dimensions["A"].width, o.column_dimensions["B"].width = 60, 16
    for r in (1, 2):
        o.cell(r, 2).number_format = PF
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(fatura: Path, anlasma: Path, provizyon: Path, police: Path, cikti: Path) -> dict:
    kal, an, pr, po = oku(fatura, anlasma, provizyon, police)
    kalem_kontrol(kal, an, pr, po)
    dos = dosya_kontrol(kal, pr, po)
    rapor_yaz(cikti, kal, dos)
    return {"kalemler": kal, "dosyalar": dos}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Fatura kalemlerini anlaşma fiyatları, provizyon onayı ve poliçe limitleriyle karşılaştırıp kesinti riskli kalemleri işaretler.")
    p.add_argument("--fatura", type=Path, default=ORNEK / "fatura_kalemleri.csv",
                   help="Dosya No, Poliçe No, Provizyon No, Hizmet Tarihi, Hizmet Kodu, Hizmet Adı, Adet, Birim Fiyat, Tutar")
    p.add_argument("--anlasma", type=Path, default=ORNEK / "anlasma_fiyatlari.csv", help="Hizmet Kodu, Hizmet Adı, Anlaşma Fiyatı, Provizyon Gerekli, Adet Sınırı")
    p.add_argument("--provizyon", type=Path, default=ORNEK / "provizyonlar.csv",
                   help="Provizyon No, Dosya No, Onay Tarihi, Geçerlilik Günü, Onaylı Tutar, Onaylı Hizmet Kodları")
    p.add_argument("--police", type=Path, default=ORNEK / "policeler.csv",
                   help="Poliçe No, Başlangıç, Bitiş, Yıllık Limit, Kullanılan Limit, Katılım Payı %%, Kapsam Dışı Kodlar")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "oss_fatura_kontrol.xlsx")
    a = p.parse_args(argv)
    for yol in (a.fatura, a.anlasma, a.provizyon, a.police):
        if not yol.exists():
            print(f"[X] Dosya bulunamadı: {yol}")
            return 1
    try:
        r = calistir(a.fatura, a.anlasma, a.provizyon, a.police, a.cikti)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    kal = r["kalemler"]
    top, risk = sum((k["tutar"] for k in kal), SIFIR), sum((k["risk"] for k in kal), SIFIR)
    print(f"[OK] {len(kal)} kalem · {len(r['dosyalar'])} dosya · toplam {tl(top)} TL · kesinti riskli {tl(risk)} TL")
    for d in r["dosyalar"]:
        for n in d["notlar"]:
            print(f"[!] {d['dosya']}: {n}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
