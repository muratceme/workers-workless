"""
Departman Bütçelerinin Konsolidasyonu — Workers / Workless kod bloğu
Finans › Bütçe ve Raporlama Uzmanı

Departmanlardan gelen bütçe şablonlarını (bir klasördeki .xlsx dosyaları) tek dosyada birleştirir, format ve toplam
hatalarını kontrol eder:
  - Şablon: üstte "Departman:" etiketi (yoksa dosya adı), başlık satırında Hesap Kodu, Gider Kalemi, Ocak … Aralık,
    Toplam (isteğe bağlı), Açıklama. Başlık satırı ilk 20 satırda aranır.
  - Hücre kontrolleri: sayı olmayan değer, metin olarak girilmiş sayı, boş ay, negatif tutar, hesaplanmamış formül.
  - Satır kontrolleri: Toplam ≠ ayların toplamı, hesap planında olmayan kod, hesap adı farkı, aynı hesabın dosyada
    tekrarı (toplanır), tamamen boş satır (atlanır).
  - Dosya kontrolleri: başlık bulunamadı, aynı departman için birden çok dosya (ada göre sonuncusu kullanılır).
  - Departman kontrolleri: bütçe tavanı aşımı, önceki yıl gerçekleşene göre değişim (--degisim eşiği).
Rapor: konsolide bütçe (hesap × ay), departman × hesap, departman özeti, uzun liste (pivot için), hatalar.
İnternete bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: 5 dosya, 4 departman, 2027
    python main.py --klasor butceler/ --hesap-plani plan.xlsx --tavanlar tavanlar.xlsx --degisim 30
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
K2 = Decimal("0.01")
SIFIR = Decimal(0)
AYLAR = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]
AY_KATLI = [("ocak", "oca"), ("subat", "sub"), ("mart", "mar"), ("nisan", "nis"), ("mayis", "may"), ("haziran", "haz"), ("temmuz", "tem"),
            ("agustos", "agu"), ("eylul", "eyl"), ("ekim", "eki"), ("kasim", "kas"), ("aralik", "ara")]


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tl(x: Decimal | None) -> str:
    return "—" if x is None else f"{x.quantize(K2, ROUND_HALF_UP):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def sayi_cevir(v) -> tuple[Decimal | None, str]:
    """(değer, durum) · durum: '' | 'metin' (metin olarak girilmiş sayı) | 'gecersiz' | 'bos'"""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None, "bos"
    if isinstance(v, bool):
        return None, "gecersiz"
    if isinstance(v, (int, float)):
        return Decimal(str(v)), ""
    s = str(v).strip().replace("TL", "").replace("₺", "").replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s), "metin"
    except InvalidOperation:
        return None, "gecersiz"


def tablo_oku(yol: Path) -> list[list]:
    icerik = None
    for kod in ("utf-8-sig", "cp1254"):
        try:
            icerik = yol.read_text(encoding=kod)
            break
        except UnicodeDecodeError:
            continue
    ilk = "\n".join(icerik.splitlines()[:10])
    return [r for r in csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)) if any(r)]


def basit_tablo(yol: Path | None) -> list[list]:
    if not yol:
        return []
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            return [list(r) for r in wb.active.iter_rows(values_only=True) if any(c not in (None, "") for c in r)]
        finally:
            wb.close()
    return tablo_oku(yol)


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Satir:
    departman: str
    dosya: str
    hucre: int
    kod: str
    ad: str
    aylar: list
    aciklama: str


@dataclass
class Dosya:
    yol: Path
    departman: str = ""
    hazirlayan: str = ""
    satirlar: list = field(default_factory=list)
    kullanildi: bool = True


def dosya_oku(yol: Path, hatalar: list) -> Dosya | None:
    d = Dosya(yol)

    def hata(onem, tur, aciklama, hucre=""):
        hatalar.append({"onem": onem, "tur": tur, "dosya": yol.name, "departman": d.departman, "hucre": hucre, "aciklama": aciklama})

    try:
        wb_v = load_workbook(yol, data_only=True)
        wb_f = load_workbook(yol, data_only=False)
    except Exception as h:  # noqa: BLE001 — bozuk dosya kullanıcıya raporlanır
        hata("Yüksek", "Dosya açılamadı", f"{type(h).__name__}: {h}")
        return None
    ws = wb_v["Bütçe"] if "Bütçe" in wb_v.sheetnames else wb_v.active
    wf = wb_f[ws.title]
    satirlar = [list(r) for r in ws.iter_rows(values_only=True)]
    baslik = None
    for i, r in enumerate(satirlar[:20]):
        k = [katla(c) for c in r]
        for j, c in enumerate(k):
            if c in ("departman", "departman adi", "birim") and j + 1 < len(r) and metin(r[j + 1]):
                d.departman = metin(r[j + 1])
            if c in ("hazirlayan", "hazirlayan kisi"):
                d.hazirlayan = metin(r[j + 1]) if j + 1 < len(r) else ""
        if "hesap kodu" in k and any(c.startswith(("ocak", "oca")) for c in k):
            baslik = i
            break
    if not d.departman:
        d.departman = yol.stem
        hata("Orta", "Departman adı yok", "Şablonda 'Departman:' bulunamadı; dosya adı kullanıldı")
    if baslik is None:
        hata("Yüksek", "Başlık bulunamadı", "İlk 20 satırda 'Hesap Kodu' ve ay sütunları bulunamadı; dosya konsolidasyona alınmadı")
        return None
    k = [katla(c) for c in satirlar[baslik]]
    sut_kod = k.index("hesap kodu")
    sut_ad = next((j for j, c in enumerate(k) if c in ("gider kalemi", "kalem", "hesap adi", "aciklama kalem")), None)
    sut_ay = []
    for tam, kisa in AY_KATLI:
        j = next((j for j, c in enumerate(k) if c in (tam, kisa) or c.startswith(tam)), None)
        sut_ay.append(j)
    eksik = [AYLAR[i] for i, j in enumerate(sut_ay) if j is None]
    if eksik:
        hata("Yüksek", "Eksik ay sütunu", f"Bulunamayan aylar: {', '.join(eksik)}; bu aylar 0 sayıldı")
    sut_top = next((j for j, c in enumerate(k) if c in ("toplam", "yillik toplam", "yil toplami")), None)
    sut_ack = next((j for j, c in enumerate(k) if c == "aciklama" and j != sut_ad), None)
    for i, r in enumerate(satirlar[baslik + 1:], baslik + 2):
        r = r + [None] * (len(k) - len(r))
        kod = metin(r[sut_kod])
        if not kod and all(r[j] in (None, "") for j in sut_ay if j is not None):
            continue
        if katla(kod).startswith(("toplam", "genel toplam")):
            continue
        if not kod:
            hata("Orta", "Hesap kodu yok", "Tutar var ama hesap kodu boş; satır alınmadı", f"A{i}")
            continue
        aylar, metin_var = [], False
        for m, j in enumerate(sut_ay):
            if j is None:
                aylar.append(SIFIR)
                continue
            v, durum = sayi_cevir(r[j])
            hucre = f"{get_column_letter(j + 1)}{i}"
            if durum == "bos":
                f = wf.cell(i, j + 1).value
                if isinstance(f, str) and f.startswith("="):
                    hata("Yüksek", "Hesaplanmamış formül", f"{AYLAR[m]}: formül değeri yok; dosyayı Excel'de açıp kaydedin", hucre)
                else:
                    hata("Bilgi", "Boş ay", f"{kod} · {AYLAR[m]} boş; 0 sayıldı", hucre)
                v = SIFIR
            elif durum == "gecersiz":
                hata("Yüksek", "Sayı değil", f"{kod} · {AYLAR[m]}: '{r[j]}' sayı değil; 0 sayıldı", hucre)
                v = SIFIR
            elif durum == "metin":
                metin_var = True
                hata("Bilgi", "Metin sayı", f"{kod} · {AYLAR[m]}: '{r[j]}' metin olarak girilmiş; {tl(v)} okundu", hucre)
            if v < 0:
                hata("Orta", "Negatif tutar", f"{kod} · {AYLAR[m]}: {tl(v)} (iade / düzeltme ise açıklama ekleyin)", hucre)
            aylar.append(v)
        if sut_top is not None:
            t, durum = sayi_cevir(r[sut_top])
            if durum == "gecersiz":
                hata("Orta", "Toplam okunamadı", f"{kod}: Toplam '{r[sut_top]}'", f"{get_column_letter(sut_top + 1)}{i}")
            elif t is not None and abs(t - sum(aylar)) > Decimal("0.5"):
                hata("Yüksek", "Toplam hatası", f"{kod}: Toplam sütunu {tl(t)}, ayların toplamı {tl(sum(aylar))} (fark {tl(t - sum(aylar))})"
                     + ("; satırda metin olarak girilmiş sayı var, Excel'in TOPLA işlevi bunları atlar" if metin_var else ""),
                     f"{get_column_letter(sut_top + 1)}{i}")
        d.satirlar.append(Satir(d.departman, yol.name, i, kod, metin(r[sut_ad]) if sut_ad is not None else "", aylar,
                                metin(r[sut_ack]) if sut_ack is not None else ""))
    if not d.satirlar:
        hata("Yüksek", "Boş bütçe", "Dosyada bütçe satırı yok")
    return d


# ----------------------------------------------------------------------------
# Konsolidasyon
# ----------------------------------------------------------------------------

def konsolide_et(klasor: Path, plan: dict[str, str], tavanlar: dict[str, tuple], degisim: Decimal) -> dict:
    hatalar = []
    dosyalar = []
    for yol in sorted(p for p in klasor.iterdir() if p.suffix.lower() in (".xlsx", ".xlsm") and not p.name.startswith("~$")):
        d = dosya_oku(yol, hatalar)
        if d:
            dosyalar.append(d)
    gruplar = defaultdict(list)
    for d in dosyalar:
        gruplar[katla(d.departman)].append(d)
    for lst in gruplar.values():
        if len(lst) > 1:
            for d in lst[:-1]:
                d.kullanildi = False
            hatalar.append({"onem": "Yüksek", "tur": "Mükerrer departman", "dosya": ", ".join(d.yol.name for d in lst), "departman": lst[-1].departman,
                            "hucre": "", "aciklama": f"Aynı departman için {len(lst)} dosya var; '{lst[-1].yol.name}' kullanıldı, diğerleri dışarıda"})
    satirlar = [s for d in dosyalar if d.kullanildi for s in d.satirlar]
    # Hesap planı ve tekrarlar
    gorulen = defaultdict(list)
    for s in satirlar:
        gorulen[(katla(s.departman), s.kod)].append(s)
        if plan:
            if s.kod not in plan:
                hatalar.append({"onem": "Yüksek", "tur": "Hesap planında yok", "dosya": s.dosya, "departman": s.departman, "hucre": f"A{s.hucre}",
                                "aciklama": f"{s.kod} '{s.ad}' hesap planında yok; doğru koda taşıyın veya plana ekleyin"})
            elif s.ad and katla(plan[s.kod]) != katla(s.ad) and katla(plan[s.kod]) not in katla(s.ad):
                hatalar.append({"onem": "Bilgi", "tur": "Hesap adı farklı", "dosya": s.dosya, "departman": s.departman, "hucre": f"B{s.hucre}",
                                "aciklama": f"{s.kod}: şablonda '{s.ad}', planda '{plan[s.kod]}'"})
    for (_, kod), lst in gorulen.items():
        if len(lst) > 1:
            hatalar.append({"onem": "Orta", "tur": "Tekrarlanan hesap", "dosya": lst[0].dosya, "departman": lst[0].departman,
                            "hucre": ", ".join(f"A{s.hucre}" for s in lst), "aciklama": f"{kod} {len(lst)} satırda var ({', '.join(s.ad for s in lst)}); toplandı"})
    # Departman özeti
    deps = defaultdict(lambda: [SIFIR] * 12)
    for s in satirlar:
        for m in range(12):
            deps[s.departman][m] += s.aylar[m]
    ozet = []
    for dep, aylar in sorted(deps.items()):
        top = sum(aylar)
        tv, onceki = tavanlar.get(katla(dep), (None, None))
        x = {"departman": dep, "toplam": top, "tavan": tv, "onceki": onceki, "fark": top - tv if tv is not None else None,
             "degisim": (top - onceki) / onceki if onceki else None, "aylar": aylar,
             "dosya": next((d.yol.name for d in dosyalar if d.kullanildi and d.departman == dep), ""),
             "hazirlayan": next((d.hazirlayan for d in dosyalar if d.kullanildi and d.departman == dep), "")}
        ozet.append(x)
        if tv is not None and top > tv:
            hatalar.append({"onem": "Yüksek", "tur": "Tavan aşımı", "dosya": x["dosya"], "departman": dep, "hucre": "",
                            "aciklama": f"Bütçe {tl(top)} TL, tavan {tl(tv)} TL (aşım {tl(top - tv)} TL)"})
        if x["degisim"] is not None and abs(x["degisim"]) > degisim:
            hatalar.append({"onem": "Orta", "tur": "Yüksek değişim", "dosya": x["dosya"], "departman": dep, "hucre": "",
                            "aciklama": f"Önceki yıl gerçekleşen {tl(onceki)} TL → bütçe {tl(top)} TL (%{x['degisim'] * 100:+.0f}); gerekçe isteyin"})
    for k, (tv, _) in tavanlar.items():
        if k not in {katla(d) for d in deps}:
            hatalar.append({"onem": "Orta", "tur": "Bütçe gelmedi", "dosya": "", "departman": k, "hucre": "", "aciklama": "Tavanı olan departmandan bütçe dosyası yok"})
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    hatalar.sort(key=lambda h: (sira[h["onem"]], h["departman"], h["tur"]))
    return {"dosyalar": dosyalar, "satirlar": satirlar, "ozet": ozet, "hatalar": hatalar, "plan": plan}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
UST = Alignment(vertical="top", wrap_text=True)
PF = "#,##0"


def _baslik(ws, basliklar, genislik=()):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "C2"


def rapor_yaz(cikti: Path, s: dict) -> None:
    wb = Workbook()
    kn = wb.active
    kn.title = "Konsolide Bütçe"
    _baslik(kn, ["Hesap Kodu", "Hesap Adı"] + AYLAR + ["Yıllık"], [11, 28] + [12] * 13)
    hes = defaultdict(lambda: [SIFIR] * 12)
    adlar = {}
    for st in s["satirlar"]:
        for m in range(12):
            hes[st.kod][m] += st.aylar[m]
        adlar.setdefault(st.kod, s["plan"].get(st.kod) or st.ad)
    for kod in sorted(hes):
        kn.append([kod, adlar[kod]] + [float(v) for v in hes[kod]] + [float(sum(hes[kod]))])
    top = [sum((hes[k][m] for k in hes), SIFIR) for m in range(12)]
    kn.append(["", "Toplam"] + [float(v) for v in top] + [float(sum(top))])
    for h in kn[kn.max_row]:
        h.font = Font(bold=True)
    for row in kn.iter_rows(min_row=2, min_col=3):
        for c in row:
            c.number_format = PF

    dh = wb.create_sheet("Departman × Hesap")
    deps = [x["departman"] for x in s["ozet"]]
    _baslik(dh, ["Hesap Kodu", "Hesap Adı"] + deps + ["Toplam"], [11, 28] + [15] * (len(deps) + 1))
    dh_t = defaultdict(lambda: defaultdict(lambda: SIFIR))
    for st in s["satirlar"]:
        dh_t[st.kod][st.departman] += sum(st.aylar)
    for kod in sorted(dh_t):
        dh.append([kod, adlar[kod]] + [float(dh_t[kod].get(d, SIFIR)) for d in deps] + [float(sum(dh_t[kod].values()))])
    for row in dh.iter_rows(min_row=2, min_col=3):
        for c in row:
            c.number_format = PF

    oz = wb.create_sheet("Departman Özeti")
    _baslik(oz, ["Departman", "Dosya", "Hazırlayan", "Yıllık Bütçe", "Tavan", "Tavan Farkı", "Önceki Yıl Gerçekleşen", "Değişim", "Hata (Yüksek)", "Onay"],
            (18, 26, 14, 15, 15, 14, 16, 9, 10, 12))
    for x in s["ozet"]:
        n = sum(1 for h in s["hatalar"] if h["departman"] == x["departman"] and h["onem"] == "Yüksek")
        oz.append([x["departman"], x["dosya"], x["hazirlayan"], float(x["toplam"]), None if x["tavan"] is None else float(x["tavan"]),
                   None if x["fark"] is None else float(x["fark"]), None if x["onceki"] is None else float(x["onceki"]),
                   None if x["degisim"] is None else float(x["degisim"]), n, ""])
        r = oz.max_row
        for j in (4, 5, 6, 7):
            oz.cell(r, j).number_format = PF
        oz.cell(r, 8).number_format = "0%"
        if x["fark"] is not None and x["fark"] > 0:
            oz.cell(r, 6).fill = PatternFill("solid", fgColor="FDE2E1")
        if n:
            oz.cell(r, 9).fill = PatternFill("solid", fgColor="FDE2E1")
        oz.cell(r, 10).fill = PatternFill("solid", fgColor="FFF4CE")
    oz.freeze_panes = "A2"
    if s["ozet"]:
        g = BarChart()
        g.title, g.height, g.width = "Bütçe ve tavan", 8, 18
        g.add_data(Reference(oz, min_col=4, max_col=5, min_row=1, max_row=1 + len(s["ozet"])), titles_from_data=True)
        g.set_categories(Reference(oz, min_col=1, min_row=2, max_row=1 + len(s["ozet"])))
        oz.add_chart(g, f"A{oz.max_row + 3}")

    ul = wb.create_sheet("Uzun Liste")
    _baslik(ul, ["Departman", "Hesap Kodu", "Hesap Adı", "Ay No", "Ay", "Tutar", "Kaynak Dosya", "Satır", "Açıklama"], (18, 11, 28, 6, 9, 13, 24, 6, 30))
    for st in s["satirlar"]:
        for m, v in enumerate(st.aylar):
            if v:
                ul.append([st.departman, st.kod, adlar[st.kod], m + 1, AYLAR[m], float(v), st.dosya, st.hucre, st.aciklama])
                ul.cell(ul.max_row, 6).number_format = PF
    ul.auto_filter.ref = f"A1:I{ul.max_row}"
    ul.freeze_panes = "A2"

    hs = wb.create_sheet("Hatalar", 0)
    _baslik(hs, ["Önem", "Tür", "Departman", "Dosya", "Hücre", "Açıklama", "Düzeltildi"], (9, 22, 18, 26, 10, 80, 11))
    hs.freeze_panes = "A2"
    for h in s["hatalar"]:
        hs.append([h["onem"], h["tur"], h["departman"], h["dosya"], h["hucre"], h["aciklama"], ""])
        hs.cell(hs.max_row, 1).fill = PatternFill("solid", fgColor=RENK[h["onem"]])
        hs.cell(hs.max_row, 6).alignment = UST
        hs.cell(hs.max_row, 7).fill = PatternFill("solid", fgColor="FFF4CE")
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(klasor: Path, cikti: Path, plan_yolu: Path | None = None, tavan_yolu: Path | None = None, degisim: Decimal = Decimal("0.30")) -> dict:
    if not klasor.is_dir():
        raise ValueError(f"Klasör bulunamadı: {klasor}")
    plan = {}
    t = basit_tablo(plan_yolu)
    if t:
        b = [katla(x) for x in t[0]]
        ik, ia = (b.index("hesap kodu") if "hesap kodu" in b else 0), (b.index("hesap adi") if "hesap adi" in b else 1)
        plan = {metin(r[ik]): metin(r[ia]) for r in t[1:] if len(r) > ik and metin(r[ik])}
    tavanlar = {}
    t = basit_tablo(tavan_yolu)
    if t:
        b = [katla(x) for x in t[0]]
        idep = next((i for i, x in enumerate(b) if x in ("departman", "birim")), 0)
        itv = next((i for i, x in enumerate(b) if "tavan" in x), None)
        ion = next((i for i, x in enumerate(b) if "onceki" in x or "gerceklesen" in x), None)
        for r in t[1:]:
            if len(r) > idep and metin(r[idep]):
                tv = sayi_cevir(r[itv])[0] if itv is not None and itv < len(r) else None
                on = sayi_cevir(r[ion])[0] if ion is not None and ion < len(r) else None
                tavanlar[katla(r[idep])] = (tv, on)
    s = konsolide_et(klasor, plan, tavanlar, degisim)
    if not s["satirlar"]:
        raise ValueError(f"{klasor}: okunabilen bütçe satırı yok")
    rapor_yaz(cikti, s)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Departman bütçe şablonlarını tek dosyada birleştirir, format ve toplam hatalarını kontrol eder.")
    p.add_argument("--klasor", type=Path, default=ORNEK / "departmanlar", help="Departman bütçe dosyalarının (.xlsx) bulunduğu klasör")
    p.add_argument("--hesap-plani", type=Path, help="İsteğe bağlı: Hesap Kodu, Hesap Adı")
    p.add_argument("--tavanlar", type=Path, help="İsteğe bağlı: Departman, Bütçe Tavanı, Önceki Yıl Gerçekleşen")
    p.add_argument("--degisim", type=float, default=30, help="Önceki yıla göre bu %%'den fazla değişim uyarı verir (varsayılan 30)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "konsolide_butce.xlsx")
    a = p.parse_args(argv)
    ornek = a.klasor == ORNEK / "departmanlar"
    plan = a.hesap_plani or (ORNEK / "hesap_plani.csv" if ornek else None)
    tavan = a.tavanlar or (ORNEK / "tavanlar.csv" if ornek else None)
    for y in (plan, tavan):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        s = calistir(a.klasor, a.cikti, plan, tavan, Decimal(str(a.degisim)) / 100)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    top = sum((x["toplam"] for x in s["ozet"]), SIFIR)
    n = sum(1 for h in s["hatalar"] if h["onem"] == "Yüksek")
    print(f"[OK] {sum(1 for d in s['dosyalar'] if d.kullanildi)} dosya · {len(s['ozet'])} departman · {len(s['satirlar'])} satır · konsolide bütçe {tl(top)} TL")
    for x in s["ozet"]:
        print(f"     {x['departman']}: {tl(x['toplam'])} TL" + (f" (tavan {tl(x['tavan'])})" if x["tavan"] is not None else ""))
    print(f"[{'!' if n else 'OK'}] {n} yüksek önemli hata · toplam {len(s['hatalar'])} bulgu (Hatalar sayfası)")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
