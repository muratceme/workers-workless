import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

ORNEK = main.ORNEK


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "k.xlsx"
        cls.s = main.calistir(ORNEK / "departmanlar", cls.cikti, ORNEK / "hesap_plani.csv", ORNEK / "tavanlar.csv")
        cls.o = {x["departman"]: x for x in cls.s["ozet"]}
        cls.h = [(h["tur"], h["departman"], h["hucre"]) for h in cls.s["hatalar"]]

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_toplamlar(self):
        self.assertEqual(self.o["Muhasebe"]["toplam"], D("8820000"))           # revize dosya: 735.000 × 12
        self.assertEqual(self.o["Pazarlama"]["toplam"], D("19605000"))
        self.assertEqual(self.o["İnsan Kaynakları"]["toplam"], D("8280000"))   # metin "15.000" okundu
        self.assertEqual(sum(x["toplam"] for x in self.s["ozet"]), D("50243000"))

    def test_hatalar(self):
        turler = {(t, d) for t, d, _ in self.h}
        self.assertTrue({("Mükerrer departman", "Muhasebe"), ("Tavan aşımı", "Pazarlama"), ("Hesap planında yok", "İnsan Kaynakları"),
                         ("Tekrarlanan hesap", "Pazarlama"), ("Negatif tutar", "Pazarlama"), ("Boş ay", "Bilgi İşlem"),
                         ("Metin sayı", "İnsan Kaynakları"), ("Yüksek değişim", "Bilgi İşlem")} <= turler)
        self.assertIn(("Toplam hatası", "İnsan Kaynakları", "O6"), self.h)
        metin = next(h for h in self.s["hatalar"] if h["hucre"] == "O7")
        self.assertIn("TOPLA", metin["aciklama"])

    def test_excel(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Hatalar", "Konsolide Bütçe", "Departman × Hesap", "Departman Özeti", "Uzun Liste"])
        kn = wb["Konsolide Bütçe"]
        self.assertEqual(kn.cell(kn.max_row, 15).value, 50243000)


class KuralTesti(unittest.TestCase):
    def test_bozuk_sablonlar(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            wb = Workbook()
            ws = wb.active
            ws.append(["Hesap Kodu", "Kalem"] + main.AYLAR)
            ws.append(["770.01", "Personel", "=1+1"] + [100] * 11)
            ws.append(["770.02", "Kira", "yok"] + [50] * 11)
            wb.save(t / "adsiz.xlsx")
            wb = Workbook()
            wb.active.append(["Bu bir şablon değil"])
            wb.save(t / "bozuk.xlsx")
            s = main.calistir(t, t / "o.xlsx")
            turler = {h["tur"] for h in s["hatalar"]}
            self.assertTrue({"Departman adı yok", "Hesaplanmamış formül", "Sayı değil", "Başlık bulunamadı"} <= turler)
            self.assertEqual(s["ozet"][0]["departman"], "adsiz")
            self.assertEqual(s["ozet"][0]["toplam"], D("1650"))
        self.assertEqual(main.sayi_cevir("1.250,50"), (D("1250.50"), "metin"))
        self.assertEqual(main.sayi_cevir(None), (None, "bos"))

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--klasor", str(Path(tmp) / "yok")]), 1)


if __name__ == "__main__":
    unittest.main()
