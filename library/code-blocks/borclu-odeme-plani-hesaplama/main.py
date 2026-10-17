"""
Borçlu Ödeme Planı Hesaplama — Gerisi Bende kod bloğu
İcra Takip › İcra Takip Elemanı (Avukat Katibi)
Toplam borç, peşinat, taksit sayısı ve faiz oranına göre faiz dahil ödeme planı ve protokol özeti çıkarır.
  Kalan anapara  = toplam borç − peşinat
  Azalan bakiye  = her taksitte eşit anapara + kalan anapara × yıllık faiz × gün ÷ 365 (gün: önceki ödeme tarihinden bu taksite)
  Eşit taksit    = aylık oran r = yıllık ÷ 12 ile annüite: taksit = A × r ÷ (1 − (1 + r)^−n); faiz = kalan × r
  Faiz oranını takip / ilam / sözleşmeye göre kullanıcı verir; 0 verilirse faizsiz plan çıkar.
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
SUTUN = {"no": ("dosya no", "icra dosya no", "dosya"), "borclu": ("borclu", "borclu adi", "ad soyad"), "borc": ("toplam borc", "borc", "dosya borcu", "tutar"),
         "pesinat": ("pesinat", "on odeme"), "taksit": ("taksit sayisi", "taksit"), "faiz": ("yillik faiz", "faiz", "faiz orani"), "ilk": ("ilk taksit tarihi", "ilk taksit", "baslangic"),
         "protokol": ("protokol tarihi", "tarih"), "yontem": ("yontem", "hesap yontemi")}


def ay_ekle(g: date, ay: int) -> date:
    y, m = divmod(g.month - 1 + ay, 12)
    yil, ayn = g.year + y, m + 1
    son = (date(yil + (ayn == 12), ayn % 12 + 1, 1) - timedelta(days=1)).day
    return date(yil, ayn, min(g.day, son))


def plan(borc: Decimal, taksit: int, faiz: Decimal, ilk: date, protokol: date, pesinat: Decimal = Decimal(0), yontem: str = "azalan", aralik: int = 1) -> dict:
    if borc <= 0 or taksit < 1 or pesinat < 0 or pesinat >= borc or faiz < 0 or ilk < protokol:
        raise ValueError("Borç 0'dan büyük, peşinat borçtan küçük, taksit en az 1, faiz 0 ya da üstü ve ilk taksit protokol tarihinden sonra olmalı.")
    ana = borc - pesinat
    satirlar, kalan, onceki = [], ana, protokol
    if yontem == "esit" and faiz > 0:
        r = faiz / 100 / 12
        tk = yuvarla(ana * r / (1 - (1 + r) ** -taksit))
    for i in range(taksit):
        vade = ay_ekle(ilk, i * aralik)
        gun = (vade - onceki).days
        if yontem == "esit" and faiz > 0:
            f = yuvarla(kalan * r)
            a = kalan if i == taksit - 1 else tk - f
        else:
            f = yuvarla(kalan * faiz / 100 * gun / 365)
            a = kalan if i == taksit - 1 else yuvarla(ana / taksit)
        kalan_once = kalan
        kalan -= a
        satirlar.append({"no": i + 1, "vade": vade, "gun": gun, "kalan_once": kalan_once, "anapara": a, "faiz": f, "taksit": a + f, "kalan": kalan})
        onceki = vade
    top_faiz = sum((s["faiz"] for s in satirlar), Decimal(0))
    return {"borc": borc, "pesinat": pesinat, "ana": ana, "faiz_orani": faiz, "yontem": yontem, "satirlar": satirlar, "toplam_faiz": top_faiz,
            "toplam_odeme": pesinat + ana + top_faiz, "son": satirlar[-1]["vade"], "protokol": protokol}


def _yontem(x) -> str | None:
    k = katla(x)
    if not k or k in ("azalan", "azalan bakiye", "esit anapara"):
        return "azalan"
    if k in ("esit", "esit taksit", "annuite", "sabit taksit"):
        return "esit"
    return None


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("borc", "taksit")):
        if not metin(r["borc"]):
            continue
        no = metin(r["no"]) or f"Satır {r['_satir']}"
        v = {a: (para(r[a]) if metin(r[a]) else Decimal(0)) for a in ("borc", "pesinat", "taksit", "faiz")}
        ilk, prot, yon = tarih(r["ilk"]), tarih(r["protokol"]) if metin(r["protokol"]) else None, _yontem(r["yontem"])
        if any(x is None for x in v.values()) or ilk is None or yon is None or (metin(r["protokol"]) and prot is None) or v["taksit"] != int(v["taksit"] or 0):
            uy.append(uyari("Yüksek", "Satır okunamadı", no, f"Satır {r['_satir']}: borç, peşinat, taksit sayısı, faiz, tarih ya da yöntem okunamadı."))
            continue
        l.append({"no": no, "borclu": metin(r["borclu"]), **v, "taksit": int(v["taksit"]), "ilk": ilk, "protokol": prot or ilk - timedelta(days=30), "protokol_yazili": prot is not None,
                  "yontem": yon})
    return l, uy


def hesapla(dosyalar: list[dict]) -> dict:
    if not dosyalar:
        raise ValueError("Hesaplanacak dosya yok.")
    sonuc, uy = [], []
    for d in dosyalar:
        try:
            p = plan(d["borc"], d["taksit"], d["faiz"], d["ilk"], d["protokol"], d["pesinat"], d["yontem"])
        except ValueError as h:
            uy.append(uyari("Yüksek", "Plan çıkarılamadı", d["no"], str(h)))
            continue
        sonuc.append({**d, **p})
        if not d["protokol_yazili"]:
            uy.append(uyari("Bilgi", "Protokol tarihi yok", d["no"], "Protokol tarihi yazılmadığı için ilk taksitten 30 gün önce kabul edildi; ilk taksit faizi buna göre hesaplandı."))
        if d["faiz"] == 0:
            uy.append(uyari("Bilgi", "Faizsiz plan", d["no"], "Faiz oranı 0; plan faizsiz çıkarıldı. Takip / ilamdaki faiz hakkından vazgeçilip vazgeçilmediğini protokolde açıkça yazın."))
        if d["taksit"] > 36:
            uy.append(uyari("Bilgi", "Uzun vade", d["no"], f"{d['taksit']} taksit; uzun vadede faiz oranı değişikliği ve tahsilat riski için ara değerlendirme tarihi koyun."))
    return {"dosyalar": sonuc, "uyarilar": uy}


def protokol_metni(s: dict) -> str:
    st = s["satirlar"]
    return (f"{s['no']} sayılı dosyada borçlu {s['borclu'] or '…'}, toplam {s['borc']:,.2f} TL borcunu {s['protokol']:%d.%m.%Y} tarihli protokolle; "
            + (f"{s['pesinat']:,.2f} TL peşinat ve " if s["pesinat"] else "")
            + f"kalan {s['ana']:,.2f} TL'yi yıllık %{s['faiz_orani']} faizle {len(st)} taksitte "
            + f"({st[0]['vade']:%d.%m.%Y} – {s['son']:%d.%m.%Y}) ödemeyi kabul eder. Toplam faiz {s['toplam_faiz']:,.2f} TL, toplam ödeme {s['toplam_odeme']:,.2f} TL'dir. "
            + "Taksitlerden birinin vadesinde ödenmemesi hâlinde uygulanacak sonuç (muacceliyet, takibe devam) protokolde ayrıca yazılmalıdır.")


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    _baslik(ws, ["Dosya", "Borçlu", "Toplam Borç", "Peşinat", "Taksitlenen", "Yıllık Faiz %", "Yöntem", "Taksit", "İlk Vade", "Son Vade", "Toplam Faiz", "Toplam Ödeme", "Protokol Özeti"],
            (14, 20, 14, 12, 14, 8, 10, 7, 11, 11, 12, 14, 90))
    for s in r["dosyalar"]:
        _satir(ws, [s["no"], s["borclu"], s["borc"], s["pesinat"], s["ana"], s["faiz_orani"], "Eşit taksit" if s["yontem"] == "esit" else "Azalan bakiye", len(s["satirlar"]),
                    s["satirlar"][0]["vade"], s["son"], s["toplam_faiz"], s["toplam_odeme"], protokol_metni(s)], para_s=(3, 4, 5, 6, 11, 12), tarih_s=(9, 10), sar=(13,))
    ws = wb.create_sheet("Ödeme Planı")
    _baslik(ws, ["Dosya", "Taksit", "Vade", "Gün", "Dönem Başı Anapara", "Anapara", "Faiz", "Taksit Tutarı", "Kalan Anapara"], (14, 7, 11, 6, 14, 13, 12, 13, 14))
    for s in r["dosyalar"]:
        if s["pesinat"]:
            _satir(ws, [s["no"], "Peşinat", s["protokol"], None, s["borc"], s["pesinat"], 0, s["pesinat"], s["ana"]], para_s=(5, 6, 7, 8, 9), tarih_s=(3,))
        for t in s["satirlar"]:
            _satir(ws, [s["no"], t["no"], t["vade"], t["gun"], t["kalan_once"], t["anapara"], t["faiz"], t["taksit"], t["kalan"]], para_s=(5, 6, 7, 8, 9), tarih_s=(3,))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(dosyalar: list[dict], cikti: Path, uy: list[dict] | None = None) -> dict:
    r = hesapla(dosyalar)
    r["uyarilar"] = sirala((uy or []) + r["uyarilar"])
    if not r["dosyalar"]:
        raise ValueError("Hiçbir dosya için plan çıkarılamadı: " + "; ".join(u["aciklama"] for u in r["uyarilar"][:3]))
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Borç, peşinat, taksit sayısı ve faiz oranına göre ödeme planı ve protokol özeti çıkarır.")
    p.add_argument("--dosyalar", type=Path, help="Dosya No, Borçlu, Toplam Borç, Peşinat, Taksit Sayısı, Yıllık Faiz %%, İlk Taksit Tarihi, Protokol Tarihi, Yöntem (Azalan / Eşit)")
    p.add_argument("--borc", help="Tek dosya: toplam borç (TL)")
    p.add_argument("--pesinat", default="0", help="Tek dosya: peşinat")
    p.add_argument("--taksit", type=int, help="Tek dosya: taksit sayısı")
    p.add_argument("--faiz", help="Tek dosya: yıllık faiz oranı, yüzde (takip / ilamdaki oran)")
    p.add_argument("--ilk", help="Tek dosya: ilk taksit tarihi GG.AA.YYYY")
    p.add_argument("--protokol", help="Tek dosya: protokol tarihi GG.AA.YYYY (varsayılan bugün)")
    p.add_argument("--yontem", choices=("azalan", "esit"), default="azalan", help="azalan = eşit anapara + kalan bakiyeye faiz (varsayılan); esit = eşit taksit")
    p.add_argument("--borclu", default="", help="Tek dosya: borçlu adı")
    p.add_argument("--dosya-no", default="Dosya", help="Tek dosya: dosya numarası")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "odeme_plani.xlsx")
    a = p.parse_args(argv)
    uy = []
    if a.borc:
        b, ps, fz, ilk = para(a.borc), para(a.pesinat), para(a.faiz) if a.faiz else None, tarih(a.ilk) if a.ilk else None
        prot = tarih(a.protokol) if a.protokol else date.today()
        if b is None or ps is None or fz is None or ilk is None or prot is None or not a.taksit or a.taksit < 1 or a.taksit > 240:
            print("[X] Tek dosyada --borc, --taksit (1-240), --faiz ve --ilk zorunlu; tarihler GG.AA.YYYY olmalı")
            return 2
        dosyalar = [{"no": a.dosya_no, "borclu": a.borclu, "borc": b, "pesinat": ps, "taksit": a.taksit, "faiz": fz, "ilk": ilk, "protokol": prot, "protokol_yazili": True, "yontem": a.yontem}]
    else:
        yol = a.dosyalar or ORNEK / "dosyalar.csv"
        if not yol.exists():
            print(f"[X] Dosya bulunamadı: {yol}")
            return 1
        try:
            dosyalar, uy = oku(yol)
        except ValueError as h:
            print(f"[X] {h}")
            return 1
    try:
        r = calistir(dosyalar, a.cikti, uy)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    for s in r["dosyalar"]:
        print(f"[OK] {s['no']}: {s['ana']:,.2f} TL · {len(s['satirlar'])} taksit · %{s['faiz_orani']} · toplam faiz {s['toplam_faiz']:,.2f} · ilk taksit {s['satirlar'][0]['taksit']:,.2f} · son vade {s['son']:%d.%m.%Y}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Plan: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
