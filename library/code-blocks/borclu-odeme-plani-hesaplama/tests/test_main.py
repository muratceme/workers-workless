import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


class PlanTesti(unittest.TestCase):
    def test_azalan(self):
        p = main.plan(D(120000), 10, D(24), date(2026, 11, 15), date(2026, 10, 14), D(20000))
        s = p["satirlar"]
        # 100.000 × %24 × 32 gün ÷ 365 = 2.104,11; ikinci taksit 90.000 × %24 × 30 ÷ 365 = 1.775,34
        self.assertEqual((s[0]["gun"], s[0]["faiz"], s[0]["taksit"], s[1]["faiz"], s[-1]["kalan"], s[-1]["vade"]), (32, D("2104.11"), D("12104.11"), D("1775.34"), D(0), date(2027, 8, 15)))
        self.assertEqual((p["ana"], p["toplam_faiz"], p["toplam_odeme"]), (D(100000), D("11059.73"), D("131059.73")))

    def test_esit(self):
        p = main.plan(D(48500), 6, D(24), date(2026, 11, 1), date(2026, 10, 14), yontem="esit")
        s = p["satirlar"]
        # aylık %2, annüite 48.500 × 0,02 ÷ (1 − 1,02^−6) = 8.658,50
        self.assertEqual([x["taksit"] for x in s[:5]], [D("8658.50")] * 5)
        self.assertEqual((s[0]["faiz"], s[-1]["kalan"], p["toplam_faiz"]), (D("970.00"), D(0), D("3451.00")))
        self.assertEqual(main.plan(D(1000), 3, D(0), date(2026, 1, 31), date(2026, 1, 1), yontem="esit")["satirlar"][1]["vade"], date(2026, 2, 28))

    def test_hatalar(self):
        for k in ({"borc": D(0)}, {"taksit": 0}, {"faiz": D(-1)}, {"pesinat": D(100)}, {"ilk": date(2026, 1, 1)}):
            args = {"borc": D(100), "taksit": 2, "faiz": D(10), "ilk": date(2026, 3, 1), "protokol": date(2026, 2, 1), "pesinat": D(0)} | k
            with self.assertRaises(ValueError):
                main.plan(**args)
        self.assertEqual([main._yontem(x) for x in ("", "Eşit taksit", "annüite", "x")], ["azalan", "esit", "esit", None])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            l, uy = main.oku(ORNEK / "dosyalar.csv")
            r = main.calistir(l, cikti, uy)
            self.assertEqual([(s["no"], s["toplam_faiz"]) for s in r["dosyalar"]],
                             [("2026/1452 E.", D("11059.73")), ("2026/1460 E.", D("3451.00")), ("2026/1477 E.", D("93821.94")), ("2026/1490 E.", D(0)), ("2026/1511 E.", D("112977.49"))])
            self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                             [("Bilgi", "Faizsiz plan", "2026/1490 E."), ("Bilgi", "Protokol tarihi yok", "2026/1511 E."), ("Bilgi", "Uzun vade", "2026/1511 E."),
                              ("Yüksek", "Plan çıkarılamadı", "2026/1502 E.")])
            self.assertIn("toplam 120,000.00 TL borcunu", main.protokol_metni(r["dosyalar"][0]))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Ödeme Planı", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "d.csv"
            y.write_text("Dosya;Borç;Taksit;Faiz;İlk Taksit\n1;10.000;4;20;01.12.2026\n2;x;4;20;01.12.2026\n3;1.000;2,5;20;01.12.2026\n4;1.000;2;20;x\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["taksit"], x["protokol"]) for x in l], len(uy)), ([("1", 4, date(2026, 11, 1))], 3))
            c = str(Path(tmp) / "x.xlsx")
            tek = ["--borc", "50.000", "--taksit", "5", "--faiz", "24", "--ilk", "01.12.2026", "--protokol", "01.11.2026", "--cikti", c]
            self.assertEqual(main.main(tek), 0)
            self.assertEqual(main.main(tek + ["--yontem", "esit", "--pesinat", "10.000"]), 0)
            self.assertEqual([main.main(v + ["--cikti", c]) for v in (["--borc", "1"], ["--borc", "1", "--taksit", "2", "--faiz", "x", "--ilk", "01.12.2026"],
                                                                       ["--borc", "1", "--taksit", "500", "--faiz", "1", "--ilk", "01.12.2026"])], [2] * 3)
            self.assertEqual(main.main(["--borc", "100", "--taksit", "2", "--faiz", "1", "--ilk", "01.01.2026", "--protokol", "01.02.2026", "--cikti", c]), 1)
            self.assertEqual(main.main(["--dosyalar", str(y), "--cikti", c]), 0)
            self.assertEqual(main.main(["--dosyalar", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Dosya;Borç;Taksit\n1;x;1\n", encoding="utf-8")
            self.assertEqual(main.main(["--dosyalar", str(y), "--cikti", c]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
