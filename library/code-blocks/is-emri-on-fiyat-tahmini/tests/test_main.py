import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
PAKET = {"bakim": {"kod": "BAKIM", "ad": "Bakım", "satirlar": [{"tur": "İşçilik", "parca": "", "miktar": D("1.5")}, {"tur": "Parça", "parca": "YAG", "miktar": D(4)},
                                                                 {"tur": "Parça", "parca": "FLT", "miktar": D(1)}]},
         "balans": {"kod": "BALANS", "ad": "", "satirlar": [{"tur": "Parça", "parca": "AGR", "miktar": D(1)}]}}
PARCA = {"yag": {"kod": "YAG", "ad": "Yağ", "birim": "lt", "fiyat": D(500)}, "flt": {"kod": "FLT", "ad": "Filtre", "birim": "adet", "fiyat": D(400)}}


def t(no, kod, adet=1):
    return {"no": no, "plaka": "P", "musteri": "", "kod": kod, "adet": D(adet)}


class HesapTesti(unittest.TestCase):
    def test_tutarlar(self):
        r = main.hesapla([t("1", "bakım"), t("2", "BAKIM", 2), t("2", "BALANS"), t("3", "YOK")], PAKET, PARCA, D(2000), D(10), D(5), D(20), D(10))
        e = {x["no"]: x for x in r["emirler"]}
        # 1: işçilik 1,5 × 2000 = 3000; parça 4 × 500 + 400 = 2400; iskonto 300 + 120; net 4980; KDV 996
        self.assertEqual((e["1"]["saat"], e["1"]["iscilik"], e["1"]["parca"], e["1"]["iskonto"], e["1"]["net"], e["1"]["kdv"], e["1"]["toplam"], e["1"]["alt"], e["1"]["ust"], e["1"]["eksik"]),
                         (D("1.5"), D(3000), D(2400), D(420), D(4980), D(996), D(5976), D("5378.40"), D("6573.60"), False))
        self.assertEqual((e["2"]["saat"], e["2"]["iscilik"], e["2"]["parca"], e["2"]["eksik"], len(e["2"]["satirlar"])), (D(3), D(6000), D(4800), True, 4))
        self.assertEqual((e["3"]["toplam"], e["3"]["eksik"]), (D(0), True))
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Bilgi", "İşçilik tanımı yok", "2 · BALANS"), ("Yüksek", "Parça fiyatı yok", "2 · AGR"), ("Yüksek", "İş kodu tanımsız", "3 · YOK")])
        with self.assertRaises(ValueError):
            main.hesapla([], PAKET, PARCA, D(1))

    def test_tur(self):
        self.assertEqual([main._tur(x) for x in ("İşçilik", "malzeme", "Yağ", "x")], ["İşçilik", "Parça", "Parça", None])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "talepler.csv", ORNEK / "is_tanimlari.csv", ORNEK / "parcalar.csv", cikti, D(2400))
            self.assertEqual([(e["no"], e["iscilik"], e["parca"], e["toplam"], e["eksik"]) for e in r["emirler"]],
                             [("IE-26-0412", D(5280), D(8315), D(16314), False), ("IE-26-0413", D(5760), D(15070), D(24996), False), ("IE-26-0414", D(3360), D(4950), D(9972), False),
                              ("IE-26-0415", D(4800), D(5065), D(11838), True), ("IE-26-0416", D(0), D(0), D(0), True)])
            self.assertEqual([u["tur"] for u in r["uyarilar"]], ["Parça fiyatı yok", "İş kodu tanımsız"])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Ön Fiyat Özeti", "Detay", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            iy, py, ty = Path(tmp) / "i.csv", Path(tmp) / "p.csv", Path(tmp) / "t.csv"
            iy.write_text("Paket Kodu;Paket Adı;Kalem Türü;Malzeme Kodu;Miktar\nA;Bakım;İşçilik;;1\nA;;Parça;X;2\nA;;Parça;;1\nA;;Boya;;1\nB;Diag;İşçilik;;0\n", encoding="utf-8")
            py.write_text("Kod;Ad;Fiyat\nX;Parça X;1.250,50\nX;Parça X;1\nY;Y;x\n", encoding="utf-8")
            ty.write_text("İş Emri;Plaka;Kod;Adet\n1;34 A 1;A;\n2;;A;0\n;;;\n", encoding="utf-8")
            p, uy = main.paket_oku(iy)
            self.assertEqual((list(p), len(p["a"]["satirlar"]), len(uy)), (["a"], 2, 3))
            pr, uy = main.parca_oku(py)
            self.assertEqual((pr["x"]["fiyat"], [u["tur"] for u in uy]), (D("1250.50"), ["Parça iki kez yazılmış", "Parça fiyatı okunamadı"]))
            tl, uy = main.talep_oku(ty)
            self.assertEqual(([(x["no"], x["adet"]) for x in tl], len(uy)), ([("1", D(1))], 1))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--talepler", str(ty), "--isler", str(iy), "--parcalar", str(py), "--cikti", c]
            self.assertEqual(main.main(arg + ["--saat-ucreti", "2.000", "--iscilik-iskonto", "10", "--parca-iskonto", "5", "--kdv", "20", "--pay", "15"]), 0)
            self.assertEqual([main.main(arg + v) for v in ([], ["--saat-ucreti", "0"], ["--saat-ucreti", "1", "--kdv", "100"], ["--saat-ucreti", "1", "--pay", "x"])], [2] * 4)
            self.assertEqual(main.main(["--talepler", str(Path(tmp) / "yok.csv"), "--saat-ucreti", "1"]), 1)
            ty.write_text("İş Emri;Kod\n;\n", encoding="utf-8")
            self.assertEqual(main.main(arg + ["--saat-ucreti", "1"]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
