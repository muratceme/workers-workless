"""
İş Emri Ön Fiyat Tahmini — Gerisi Bende kod bloğu
Satış Sonrası Hizmetler (Servis) › Servis Danışmanı
İş türü tanımları (işçilik saati + parça listesi), parça fiyat listesi ve saat ücretinden iş emri bazında ön fiyat tahmini çıkarır.
  İşçilik  = iş tanımındaki saat × adet × --saat-ucreti
  Parça    = iş tanımındaki miktar × adet × parça birim fiyatı (KDV hariç)
  İskonto  = işçilik × --iscilik-iskonto % + parça × --parca-iskonto %
  KDV      = (ara toplam − iskonto) × --kdv %;  toplam = net + KDV
  Aralık   = toplam × (1 ± --pay %); son fiyat araç incelendikten sonra kesinleşir
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
PAKET = {"kod": ("is kodu", "paket kodu", "kod"), "ad": ("is adi", "paket adi", "aciklama", "is"), "tur": ("tur", "kalem turu", "tip"), "parca": ("parca kodu", "parca", "malzeme kodu"),
         "miktar": ("miktar", "adet", "saat")}
PARCA = {"kod": ("parca kodu", "kod", "malzeme kodu"), "ad": ("parca adi", "ad", "aciklama", "malzeme adi"), "birim": ("birim",), "fiyat": ("birim fiyat", "fiyat", "liste fiyati", "kdv haric fiyat")}
TALEP = {"no": ("is emri no", "is emri", "no", "talep no"), "plaka": ("plaka",), "musteri": ("musteri", "ad soyad"), "kod": ("is kodu", "paket kodu", "kod"), "adet": ("adet", "miktar")}
SIFIR = Decimal(0)


def _tur(x) -> str | None:
    k = katla(x)
    if k in ("iscilik", "isc", "labor", "saat", "i"):
        return "İşçilik"
    if k in ("parca", "malzeme", "yedek parca", "sarf", "yag", "p"):
        return "Parça"
    return None


def paket_oku(yol: Path) -> tuple[dict, list[dict]]:
    uy, d = [], {}
    for r in kayitlar(yol, PAKET, ("kod", "tur", "miktar")):
        kod = metin(r["kod"])
        if not kod:
            continue
        tur, mik = _tur(r["tur"]), para(r["miktar"])
        if tur is None or mik is None or mik <= 0 or (tur == "Parça" and not metin(r["parca"])):
            uy.append(uyari("Yüksek", "İş tanımı satırı okunamadı", kod, f"Satır {r['_satir']}: tür ('{metin(r['tur'])}' — İşçilik / Parça), miktar ya da parça kodu okunamadı."))
            continue
        p = d.setdefault(katla(kod), {"kod": kod, "ad": metin(r["ad"]), "satirlar": []})
        p["ad"] = p["ad"] or metin(r["ad"])
        p["satirlar"].append({"tur": tur, "parca": metin(r["parca"]), "miktar": mik})
    return d, uy


def parca_oku(yol: Path) -> tuple[dict, list[dict]]:
    uy, d = [], {}
    for r in kayitlar(yol, PARCA, ("kod", "fiyat")):
        kod = metin(r["kod"])
        if not kod:
            continue
        f = para(r["fiyat"])
        if f is None or f < 0:
            uy.append(uyari("Yüksek", "Parça fiyatı okunamadı", kod, f"Satır {r['_satir']}: birim fiyat okunamadı."))
            continue
        if katla(kod) in d:
            uy.append(uyari("Orta", "Parça iki kez yazılmış", kod, f"Satır {r['_satir']}: ikinci fiyat alınmadı."))
            continue
        d[katla(kod)] = {"kod": kod, "ad": metin(r["ad"]), "birim": metin(r["birim"]) or "adet", "fiyat": f}
    return d, uy


def talep_oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, TALEP, ("no", "kod")):
        no, kod = metin(r["no"]), metin(r["kod"])
        if not no and not kod:
            continue
        adet = para(r["adet"]) if metin(r["adet"]) else Decimal(1)
        if not no or not kod or adet is None or adet <= 0:
            uy.append(uyari("Yüksek", "Talep satırı okunamadı", no or kod, f"Satır {r['_satir']}: iş emri no, iş kodu ya da adet okunamadı."))
            continue
        l.append({"no": no, "plaka": metin(r["plaka"]), "musteri": metin(r["musteri"]), "kod": kod, "adet": adet})
    return l, uy


def hesapla(talepler: list[dict], paketler: dict, parcalar: dict, saat_ucreti: Decimal, isc_isk: Decimal = SIFIR, prc_isk: Decimal = SIFIR, kdv: Decimal = Decimal(20),
            pay: Decimal = Decimal(10)) -> dict:
    if not talepler:
        raise ValueError("Talep satırı yok.")
    uy, emirler = [], {}
    for t in talepler:
        e = emirler.setdefault(t["no"], {"no": t["no"], "plaka": t["plaka"], "musteri": t["musteri"], "satirlar": [], "eksik": False})
        e["plaka"], e["musteri"] = e["plaka"] or t["plaka"], e["musteri"] or t["musteri"]
        p = paketler.get(katla(t["kod"]))
        if p is None:
            uy.append(uyari("Yüksek", "İş kodu tanımsız", f"{t['no']} · {t['kod']}", "İş tanımları listesinde yok; bu iş fiyata girmedi."))
            e["eksik"] = True
            continue
        if not any(s["tur"] == "İşçilik" for s in p["satirlar"]):
            uy.append(uyari("Bilgi", "İşçilik tanımı yok", f"{t['no']} · {p['kod']}", "İş tanımında işçilik satırı yok; yalnız parça fiyatlandı."))
        for s in p["satirlar"]:
            mik = s["miktar"] * t["adet"]
            if s["tur"] == "İşçilik":
                e["satirlar"].append({"is": p["ad"] or p["kod"], "tur": "İşçilik", "kod": "", "ad": f"İşçilik ({p['ad'] or p['kod']})", "birim": "saat", "miktar": mik, "fiyat": saat_ucreti,
                                      "tutar": yuvarla(mik * saat_ucreti), "fiyat_yok": False})
                continue
            pr = parcalar.get(katla(s["parca"]))
            if pr is None:
                uy.append(uyari("Yüksek", "Parça fiyatı yok", f"{t['no']} · {s['parca']}", f"{p['ad'] or p['kod']}: parça fiyat listesinde yok; tutar 0 yazıldı, tahmin eksik."))
                e["eksik"] = True
            e["satirlar"].append({"is": p["ad"] or p["kod"], "tur": "Parça", "kod": s["parca"], "ad": pr["ad"] if pr else "(fiyat listesinde yok)", "birim": pr["birim"] if pr else "",
                                  "miktar": mik, "fiyat": pr["fiyat"] if pr else None, "tutar": yuvarla(mik * pr["fiyat"]) if pr else SIFIR, "fiyat_yok": pr is None})
    for e in emirler.values():
        isc = sum((s["tutar"] for s in e["satirlar"] if s["tur"] == "İşçilik"), SIFIR)
        prc = sum((s["tutar"] for s in e["satirlar"] if s["tur"] == "Parça"), SIFIR)
        isk = yuvarla(isc * isc_isk / 100) + yuvarla(prc * prc_isk / 100)
        net = isc + prc - isk
        k = yuvarla(net * kdv / 100)
        top = net + k
        e.update({"iscilik": isc, "parca": prc, "saat": sum((s["miktar"] for s in e["satirlar"] if s["tur"] == "İşçilik"), SIFIR), "iskonto": isk, "net": net, "kdv": k, "toplam": top,
                  "alt": yuvarla(top * (100 - pay) / 100), "ust": yuvarla(top * (100 + pay) / 100)})
    return {"emirler": list(emirler.values()), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, ayar: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ön Fiyat Özeti"
    _baslik(ws, ["İş Emri", "Plaka", "Müşteri", "İşçilik Saati", "İşçilik", "Parça", "İskonto", "Net (KDV hariç)", "KDV", "Toplam (KDV dahil)", "Alt Sınır", "Üst Sınır", "Eksik Fiyat"],
            (12, 12, 20, 9, 12, 12, 11, 13, 11, 14, 12, 12, 9))
    for e in r["emirler"]:
        _satir(ws, [e["no"], e["plaka"], e["musteri"], e["saat"], e["iscilik"], e["parca"], e["iskonto"], e["net"], e["kdv"], e["toplam"], e["alt"], e["ust"], "Evet" if e["eksik"] else ""],
               para_s=(4, 5, 6, 7, 8, 9, 10, 11, 12), dolgu={13: KIRMIZI if e["eksik"] else None})
    ws.append([])
    ws.append([f"Saat ücreti {ayar['saat']:,.2f} · işçilik iskontosu %{ayar['isc']} · parça iskontosu %{ayar['prc']} · KDV %{ayar['kdv']} · tahmin aralığı ±%{ayar['pay']}."])
    ws.append(["Ön fiyat tahminidir; araç incelendikten sonra ek iş ya da parça gerekirse müşteri onayı alınarak güncellenir."])
    ws = wb.create_sheet("Detay")
    _baslik(ws, ["İş Emri", "İş", "Tür", "Parça Kodu", "Açıklama", "Birim", "Miktar", "Birim Fiyat", "Tutar"], (12, 26, 9, 14, 30, 7, 8, 12, 12))
    for e in r["emirler"]:
        for s in e["satirlar"]:
            _satir(ws, [e["no"], s["is"], s["tur"], s["kod"], s["ad"], s["birim"], s["miktar"], s["fiyat"], s["tutar"]], para_s=(7, 8, 9), dolgu={5: KIRMIZI if s["fiyat_yok"] else None})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(talep_yol: Path, paket_yol: Path, parca_yol: Path, cikti: Path, saat_ucreti: Decimal, isc_isk: Decimal = SIFIR, prc_isk: Decimal = SIFIR, kdv: Decimal = Decimal(20),
             pay: Decimal = Decimal(10)) -> dict:
    paketler, uy1 = paket_oku(paket_yol)
    parcalar, uy2 = parca_oku(parca_yol)
    talepler, uy3 = talep_oku(talep_yol)
    r = hesapla(talepler, paketler, parcalar, saat_ucreti, isc_isk, prc_isk, kdv, pay)
    r["uyarilar"] = sirala(uy1 + uy2 + uy3 + r["uyarilar"])
    rapor_yaz(cikti, r, {"saat": saat_ucreti, "isc": isc_isk, "prc": prc_isk, "kdv": kdv, "pay": pay})
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="İş emri talepleri için işçilik ve parça ön fiyat tahmini çıkarır.")
    p.add_argument("--talepler", type=Path, help="İş Emri No, Plaka, Müşteri, İş Kodu, Adet")
    p.add_argument("--isler", type=Path, help="İş tanımları: İş Kodu, İş Adı, Tür (İşçilik / Parça), Parça Kodu, Miktar (saat ya da adet)")
    p.add_argument("--parcalar", type=Path, help="Parça Kodu, Parça Adı, Birim, Birim Fiyat (KDV hariç)")
    p.add_argument("--saat-ucreti", help="İşçilik saat ücreti, KDV hariç (örnek veride 2400)")
    p.add_argument("--iscilik-iskonto", default="0", help="İşçilik iskonto yüzdesi (varsayılan 0)")
    p.add_argument("--parca-iskonto", default="0", help="Parça iskonto yüzdesi (varsayılan 0)")
    p.add_argument("--kdv", default="20", help="KDV oranı yüzdesi (varsayılan 20)")
    p.add_argument("--pay", default="10", help="Tahmin aralığı ± yüzde (varsayılan 10)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "is_emri_on_fiyat.xlsx")
    a = p.parse_args(argv)
    ornek = a.talepler is None
    yollar = [a.talepler or ORNEK / "talepler.csv", a.isler or ORNEK / "is_tanimlari.csv", a.parcalar or ORNEK / "parcalar.csv"]
    saat = para(a.saat_ucreti) if a.saat_ucreti else (Decimal(2400) if ornek else None)
    yz = [para(x) for x in (a.iscilik_iskonto, a.parca_iskonto, a.kdv, a.pay)]
    if saat is None or saat <= 0 or any(v is None or not 0 <= v < 100 for v in yz):
        print("[X] --saat-ucreti zorunlu ve 0'dan büyük; iskonto, KDV ve pay 0-99 arası olmalı")
        return 2
    for y in yollar:
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(*yollar, a.cikti, saat, *yz)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    for e in r["emirler"]:
        print(f"[OK] {e['no']} {e['plaka']}: işçilik {e['iscilik']:,.2f} + parça {e['parca']:,.2f} → toplam {e['toplam']:,.2f} (aralık {e['alt']:,.2f} – {e['ust']:,.2f})"
              + (" · EKSİK FİYAT" if e["eksik"] else ""))
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
