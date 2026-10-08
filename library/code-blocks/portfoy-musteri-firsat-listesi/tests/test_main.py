import json
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))

import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402


def calistir(**kw):
    with tempfile.TemporaryDirectory() as tmp:
        s = main.calistir(main.ORNEK, Path(tmp) / "p.xlsx", bugun=date(2026, 10, 8), **kw)
        s["wb"] = load_workbook(Path(tmp) / "p.xlsx")
    s["m"] = {m.no: m for m in s["musteriler"]}
    return s


class OrnekTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.s = calistir()
        cls.m = cls.s["m"]

    def urunler(self, no):
        return [u for u, _, _ in self.m[no].firsatlar]

    def test_firsatlar(self):
        self.assertEqual([(u, p) for u, p, _ in self.m["B1001"].firsatlar], [("kart", 40), ("bes", 35), ("talimat", 15)])
        self.assertEqual(self.urunler("B1003"), ["vade", "bes"])                  # vadesi 7 gün sonra
        self.assertEqual(self.m["B1002"].firsatlar[0][:2], ("mevduat", 40))      # 640.000 TL vadesiz, yüksek bakiye bonusu
        self.assertEqual(self.urunler("B1005"), ["reaktivasyon", "mobil"])        # 110 gün işlem yok
        self.assertEqual(self.urunler("B1012")[:2], ["bes", "kasko"])
        self.assertEqual(self.urunler("B1011"), [])                               # gelir eşiklerin altında

    def test_sorumlu_satis(self):
        self.assertNotIn("kart", self.urunler("B1007"))                           # gecikmesi var
        self.assertNotIn("kredi_bitis", self.urunler("B1007"))
        self.assertNotIn("kredi_bitis", self.urunler("B1004"))                    # limit kullanımı %94
        self.assertIn("Kart limit kullanımı %94", self.m["B1004"].notlar[0])
        self.assertIn("zorunlu ürün olarak sunulamaz", dict((u, g) for u, _, g in self.m["B1002"].firsatlar)["hayat"])

    def test_izin(self):
        self.assertIn("izni yok", self.m["B1008"].kanal)
        s = calistir(izinsiz_haric=True)
        self.assertNotIn("B1008", s["m"])

    def test_urun_sayisi_ve_excel(self):
        self.assertEqual(main.urun_sayisi(self.m["B1010"]), 8)
        wb = self.s["wb"]
        self.assertEqual(wb.sheetnames, ["Fırsat Listesi", "Müşteri Özeti", "Ürün Penetrasyonu", "Kurallar"])
        f = wb["Fırsat Listesi"]
        self.assertEqual((f["B2"].value, f["D2"].value, f["E2"].value), ("B1001", "Kredi kartı", 40))
        self.assertEqual(f.max_row, 1 + 24)

    def test_ayar_dosyasi(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "a.json"
            y.write_text(json.dumps({"min_gelir_kart": 70000, "puan": {"kart": 5}}), encoding="utf-8")
            s = main.calistir(main.ORNEK, Path(tmp) / "p.xlsx", y, date(2026, 10, 8))
        m = {x.no: x for x in s["musteriler"]}
        self.assertNotIn("kart", [u for u, _, _ in m["B1001"].firsatlar])
        self.assertEqual(s["ayar"]["puan"]["bes"], 25)

    def test_evet(self):
        self.assertEqual([main.evet(x) for x in ("E", "evet", "H", "", "3", "0", "Var")], [True, True, False, False, True, False, True])


if __name__ == "__main__":
    unittest.main()
