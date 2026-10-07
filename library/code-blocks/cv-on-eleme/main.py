"""
CV Ön Eleme — Workers / Workless kod bloğu
İnsan Kaynakları › İnsan Kaynakları Uzman Yardımcısı

CV'leri (.pdf, .docx, .txt) ilanın ZORUNLU kriterlerine göre ilk filtreden geçirir ve her aday için
GEÇTİ / ELENDİ / MANUEL KONTROL kararını kriter bazında gerekçesiyle listeler:
  - en az deneyim yılı, en az eğitim seviyesi, zorunlu yabancı diller, zorunlu beceriler,
    "en az N tanesi" beceri grupları ve zorunlu ifadeler (sertifika, belge vb.)
CV'den okunamayan bilgi adayı elemez, "manuel kontrol"e düşürür (taranmış PDF, tarihsiz deneyim vb.).
Kriterler ortak ayrımcılık tarayıcısından geçer: yaş, cinsiyet, medeni hâl, görünüş, sağlık, din, köken gibi
kriterler kabul edilmez; ehliyet, askerlik, uyruk gibi "dikkat" kriterleri ancak işe dayalı gerekçesiyle
kullanılabilir. Ad, fotoğraf ve iletişim bilgisi karar için kullanılmaz. İnternete bağlanmaz.

CV okuma çekirdeği: CV Raporlama kod bloğu (cv_cekirdek.py).

Kullanım:
    python main.py                                           # örnek CV'ler ve kriterlerle dener
    python main.py --girdi ./cvler --kriterler kriterler.json
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import ayrimcilik
import cv_cekirdek as cek

BURASI = Path(__file__).resolve().parent
SEVIYE = {ad: sira for ad, sira, _ in cek.EGITIM_SEVIYELERI}
DESTEKLENEN = {".pdf", ".docx", ".txt"}


def kriterleri_oku(yol: Path) -> dict:
    k = json.loads(yol.read_text(encoding="utf-8"))
    k = {a: v for a, v in k.items() if not a.startswith("_")}
    if k.get("en_az_egitim") and k["en_az_egitim"] not in SEVIYE:
        raise SystemExit(f"en_az_egitim şunlardan biri olmalı: {', '.join(SEVIYE)}")
    bilinmeyen = set(k.get("zorunlu_diller", [])) - set(cek.DILLER)
    if bilinmeyen:
        raise SystemExit(f"Tanınmayan dil: {', '.join(bilinmeyen)} (tanınanlar: {', '.join(cek.DILLER)})")
    return k


def kriter_denetimi(k: dict) -> list[str]:
    """Ayrımcı kriterleri reddeder; 'dikkat' kriterleri gerekçe ister. Dönüş: hata listesi."""
    hatalar = []
    metinler = [json.dumps({a: v for a, v in k.items() if a != "zorunlu_ifadeler"}, ensure_ascii=False)]
    for z in k.get("zorunlu_ifadeler", []):
        ifade = z["ifade"] if isinstance(z, dict) else str(z)
        gerekce = z.get("gerekce", "") if isinstance(z, dict) else ""
        for b in ayrimcilik.tara(ifade):
            if b["seviye"] == "yüksek":
                hatalar.append(f"'{ifade}': {b['aciklama']} — ön elemede kullanılamaz")
            elif not gerekce.strip():
                hatalar.append(f"'{ifade}': {b['aciklama']} — yalnız işe dayalı gerekçeyle kullanılabilir ('gerekce' alanı ekleyin)")
    for m in metinler:
        for b in ayrimcilik.tara(m):
            if b["seviye"] == "yüksek":
                hatalar.append(f"Kriterlerde '{b['ifade']}': {b['aciklama']} — ön elemede kullanılamaz")
    return hatalar


def degerlendir(a: cek.Aday, metin_sade: str, k: dict, okunamadi: bool) -> tuple[str, list[tuple[str, str, str]]]:
    """Dönüş: (karar, [(kriter, sonuç ✓/✗/?, açıklama)])"""
    sonuc = []
    if okunamadi:
        return "MANUEL KONTROL", [("CV okunabilirliği", "?", "; ".join(a.uyarilar) or "Metin çıkarılamadı")]
    if k.get("en_az_deneyim_yil"):
        if a.deneyim_yil == 0:
            sonuc.append(("Deneyim", "?", "Tarih aralığı bulunamadı; deneyim CV'den okunamadı"))
        else:
            ok = a.deneyim_yil >= k["en_az_deneyim_yil"]
            sonuc.append(("Deneyim", "✓" if ok else "✗", f"{a.deneyim_yil:g} yıl (en az {k['en_az_deneyim_yil']:g})"))
    if k.get("en_az_egitim"):
        if not a.egitim:
            sonuc.append(("Eğitim", "?", "Eğitim seviyesi okunamadı"))
        else:
            ok = SEVIYE[a.egitim] >= SEVIYE[k["en_az_egitim"]]
            sonuc.append(("Eğitim", "✓" if ok else "✗", f"{a.egitim} (en az {k['en_az_egitim']})"))
    for d in k.get("zorunlu_diller", []):
        sonuc.append((f"Dil: {d}", "✓" if d in a.diller else "✗", "CV'de geçiyor" if d in a.diller else "CV'de geçmiyor"))
    for b in k.get("zorunlu_beceriler", []):
        var = cek.kelime_var(metin_sade, b)
        sonuc.append((f"Beceri: {b}", "✓" if var else "✗", "CV'de geçiyor" if var else "CV'de geçmiyor"))
    for g in k.get("beceri_gruplari", []):
        bulunan = [b for b in g["liste"] if cek.kelime_var(metin_sade, b)]
        ok = len(bulunan) >= g.get("en_az", 1)
        sonuc.append((f"En az {g.get('en_az', 1)}: {', '.join(g['liste'])}", "✓" if ok else "✗", ", ".join(bulunan) or "hiçbiri yok"))
    for z in k.get("zorunlu_ifadeler", []):
        ifade = z["ifade"] if isinstance(z, dict) else str(z)
        var = cek.kelime_var(metin_sade, ifade)
        sonuc.append((f"İfade: {ifade}", "✓" if var else "✗", "CV'de geçiyor" if var else "CV'de geçmiyor"))
    if any(s == "✗" for _, s, _ in sonuc):
        return "ELENDİ", sonuc
    if any(s == "?" for _, s, _ in sonuc):
        return "MANUEL KONTROL", sonuc
    return "GEÇTİ", sonuc


def calistir(girdi: Path, kriter_yolu: Path, cikti: Path) -> dict:
    k = kriterleri_oku(kriter_yolu)
    hatalar = kriter_denetimi(k)
    if hatalar:
        raise SystemExit("Kriterler ayrımcılık denetiminden geçmedi:\n  - " + "\n  - ".join(hatalar))
    dosyalar = sorted(p for p in girdi.iterdir() if p.suffix.lower() in DESTEKLENEN) if girdi.is_dir() else [girdi]
    beceriler = list(dict.fromkeys(k.get("zorunlu_beceriler", []) + [b for g in k.get("beceri_gruplari", []) for b in g["liste"]]))
    adaylar = []
    for yol in dosyalar:
        a = cek.cv_analiz(yol, beceriler)
        okunamadi = any(u.startswith(("Okunamadı", "Metin çıkarılamadı")) for u in a.uyarilar)
        try:
            ms = cek.sade(cek.metin_oku(yol)) if not okunamadi else ""
        except Exception:  # noqa: BLE001 — bozuk dosya: manuel kontrole düşer
            ms, okunamadi = "", True
        karar, sonuc = degerlendir(a, ms, k, okunamadi)
        adaylar.append({"aday": a, "karar": karar, "sonuc": sonuc})
    _rapor(adaylar, k, cikti)
    return {"adaylar": adaylar, "kriterler": k}


BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
KARAR_DOLGU = {"GEÇTİ": PatternFill("solid", fgColor="E3F5E1"), "ELENDİ": PatternFill("solid", fgColor="FDE2E1"),
               "MANUEL KONTROL": PatternFill("solid", fgColor="FFF4CE")}
UST = Alignment(vertical="top", wrap_text=True)


def _rapor(adaylar, k, cikti):
    wb = Workbook()
    ws = wb.active
    ws.title = "Ön Eleme"
    kriterler = list(dict.fromkeys(ad for x in adaylar for ad, _, _ in x["sonuc"]))
    ws.append(["Dosya", "Ad Soyad", "E-posta", "Telefon", "Karar", "Elenme / Kontrol Gerekçesi"] + kriterler + ["İK Onayı"])
    for c in ws[1]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI
    sira = {"GEÇTİ": 0, "MANUEL KONTROL": 1, "ELENDİ": 2}
    for x in sorted(adaylar, key=lambda x: (sira[x["karar"]], -x["aday"].deneyim_yil)):
        a, s = x["aday"], {ad: (sonuc, acik) for ad, sonuc, acik in x["sonuc"]}
        gerekce = "; ".join(f"{ad}: {acik}" for ad, (sonuc, acik) in s.items() if sonuc != "✓")
        ws.append([a.dosya, a.ad_soyad, a.eposta, a.telefon, x["karar"], gerekce]
                  + [f"{s[kr][0]} {s[kr][1]}" if kr in s else "" for kr in kriterler] + [""])
        ws.cell(ws.max_row, 5).fill = KARAR_DOLGU[x["karar"]]
        ws.cell(ws.max_row, 7 + len(kriterler)).fill = KARAR_DOLGU["MANUEL KONTROL"]
        for c in ws[ws.max_row]:
            c.alignment = UST
    for j, w in enumerate([22, 20, 24, 16, 15, 50] + [24] * len(kriterler) + [10], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    kr = wb.create_sheet("Kriterler")
    kr.append(["Kriter", "Değer"])
    for c in kr[1]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI
    for a, v in k.items():
        kr.append([a, json.dumps(v, ensure_ascii=False) if not isinstance(v, (int, float, str)) else v])
    kr.column_dimensions["A"].width = 24
    kr.column_dimensions["B"].width = 100

    b = wb.create_sheet("Bilgi")
    from collections import Counter
    say = Counter(x["karar"] for x in adaylar)
    for s in [["Sonuç", f"Geçti {say['GEÇTİ']} · Manuel kontrol {say['MANUEL KONTROL']} · Elendi {say['ELENDİ']}"],
              ["Karar kuralı", "Herhangi bir zorunlu kriter ✗ ise ELENDİ; ✗ yok ama okunamayan (?) varsa MANUEL KONTROL; hepsi ✓ ise GEÇTİ"],
              ["Adil değerlendirme", "Yalnız işe dayalı kriterler kullanılır; ad, fotoğraf, iletişim bilgisi karara girmez. Kriterler ayrımcılık "
                                     "denetiminden geçmiştir (4857 s. Kanun md. 5; 6701 s. Kanun md. 3, 6, 7)."],
              ["Sınırlama", "Kelime ve tarih tabanlı okuma yapılır; eş anlamlı ifadeler, tablo/görsel içerikli CV'ler kaçabilir. "
                            "Elenen adayları göndermeden önce örneklem kontrolü yapın."],
              ["KVKK", "CV'ler kişisel veri içerir; rapor yetkisiz kişilerle paylaşılmamalı ve saklama süresine uyulmalıdır."],
              ["Oluşturulma", datetime.now().strftime("%d.%m.%Y %H:%M")]]:
        b.append(s)
    b.column_dimensions["A"].width = 20
    b.column_dimensions["B"].width = 120
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description="CV'leri ilanın zorunlu kriterlerine göre ilk filtreden geçirir.")
    ap.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "cvler", help="CV klasörü veya tek CV (.pdf, .docx, .txt)")
    ap.add_argument("--kriterler", type=Path, default=BURASI / "kriterler_ornek.json", help="Zorunlu kriterler (JSON)")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "cv_on_eleme.xlsx")
    a = ap.parse_args(argv)
    s = calistir(a.girdi, a.kriterler, a.cikti)
    for x in s["adaylar"]:
        isaret = {"GEÇTİ": "[OK]", "ELENDİ": "[X] ", "MANUEL KONTROL": "[?] "}[x["karar"]]
        neden = "; ".join(f"{ad}: {acik}" for ad, sonuc, acik in x["sonuc"] if sonuc != "✓")
        print(f"{isaret} {x['aday'].dosya:<24} {x['karar']:<15} {neden}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
