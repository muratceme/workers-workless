"""
CV Raporlama — Workers / Workless kod bloğu
İnsan Kaynakları › İK Uzmanı

Bir klasördeki CV'leri (PDF, DOCX, TXT) okur; iletişim bilgisi, toplam deneyim,
en yüksek eğitim, yabancı dil ve beceri alanlarını kural tabanlı olarak çıkarır
ve tek bir Excel raporu üretir. İsteğe bağlı iş ilanı verilirse ilan uyum yüzdesi
hesaplar. İnternete bağlanmaz, veriniz bilgisayarınızdan çıkmaz.

Kullanım:
    python main.py                                   # örnek veriyle dener
    python main.py --girdi ./cvler --cikti cikti/aday_raporu.xlsx
    python main.py --girdi ./cvler --ilan ilan.txt --beceriler beceriler.txt
"""
from __future__ import annotations

import argparse
import logging
import re
import sys
import zipfile
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
DESTEKLENEN = {".pdf", ".docx", ".txt"}

# ----------------------------------------------------------------------------
# Metin okuma
# ----------------------------------------------------------------------------

def pdf_oku(yol: Path) -> str:
    from pypdf import PdfReader

    logging.getLogger("pypdf").setLevel(logging.ERROR)  # bozuk dosya uyarılarını rapora bırak
    return "\n".join((sayfa.extract_text() or "") for sayfa in PdfReader(str(yol)).pages)


def docx_oku(yol: Path) -> str:
    """DOCX dosyasını ek kütüphane olmadan okur (DOCX bir ZIP + XML paketidir)."""
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(yol) as z:
        kok = ElementTree.fromstring(z.read("word/document.xml"))
    satirlar = []
    for p in kok.iter(f"{ns}p"):
        satirlar.append("".join(t.text or "" for t in p.iter(f"{ns}t")))
    return "\n".join(satirlar)


def txt_oku(yol: Path) -> str:
    for kodlama in ("utf-8-sig", "cp1254", "latin-1"):
        try:
            return yol.read_text(encoding=kodlama)
        except UnicodeDecodeError:
            continue
    return ""


def metin_oku(yol: Path) -> str:
    uzanti = yol.suffix.lower()
    if uzanti == ".pdf":
        return pdf_oku(yol)
    if uzanti == ".docx":
        return docx_oku(yol)
    return txt_oku(yol)


# ----------------------------------------------------------------------------
# Yardımcılar
# ----------------------------------------------------------------------------

_TR_KUCUK = str.maketrans({"I": "ı", "İ": "i"})
_TR_ASCII = str.maketrans("çğıöşüâîû", "cgiosuaiu")


def kucuk(s: str) -> str:
    """Türkçe kurallarına göre küçük harf (I→ı, İ→i)."""
    return s.translate(_TR_KUCUK).lower()


def sade(s: str) -> str:
    """Karşılaştırma için: küçük harf + Türkçe karakterleri ASCII'ye indir."""
    return kucuk(s).translate(_TR_ASCII)


def kelime_var(metin_sade: str, ifade: str) -> bool:
    ifade_sade = re.escape(sade(ifade))
    return re.search(rf"(?<![a-z0-9]){ifade_sade}(?![a-z0-9])", metin_sade) is not None


# ----------------------------------------------------------------------------
# Alan çıkarımı
# ----------------------------------------------------------------------------

EPOSTA = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
TELEFON = re.compile(r"(?:\+?90[\s-]?)?\(?0?(5\d{2})\)?[\s.-]?(\d{3})[\s.-]?(\d{2})[\s.-]?(\d{2})")
LINKEDIN = re.compile(r"(?:https?://)?(?:www\.)?linkedin\.com/in/[\w\-%]+/?", re.I)

AY = r"(?:0?[1-9]|1[0-2])"
YIL = r"(?:19[6-9]\d|20[0-4]\d)"
SIMDI = r"(?:günümüz|gunumuz|halen|devam|hâlâ|hala|present|current|now|şu an|su an)"
ARALIK = re.compile(
    rf"(?:{AY}[./])?({YIL})\s*(?:-|–|—|to|ile|/)\s*(?:(?:{AY}[./])?({YIL})|({SIMDI}))",
    re.I,
)

EGITIM_SEVIYELERI = [  # (seviye, sıra, anahtar ifadeler) — yüksekten düşüğe
    ("Doktora", 5, ["doktora", "phd", "ph.d"]),
    ("Yüksek Lisans", 4, ["yüksek lisans", "master", "mba", "msc", "m.sc"]),
    ("Lisans", 3, ["lisans", "bachelor", "üniversitesi", "fakültesi", "bsc", "b.sc"]),
    ("Ön Lisans", 2, ["ön lisans", "önlisans", "meslek yüksekokulu", "myo", "associate"]),
    ("Lise", 1, ["lise", "high school", "anadolu lisesi"]),
]

DILLER = {
    "İngilizce": ["ingilizce", "english"],
    "Almanca": ["almanca", "german", "deutsch"],
    "Fransızca": ["fransızca", "french"],
    "İspanyolca": ["ispanyolca", "spanish"],
    "İtalyanca": ["italyanca", "italian"],
    "Rusça": ["rusça", "russian"],
    "Arapça": ["arapça", "arabic"],
    "Çince": ["çince", "chinese", "mandarin"],
    "Japonca": ["japonca", "japanese"],
}

BASLIK_KELIMELERI = {"cv", "özgeçmiş", "ozgecmis", "curriculum", "vitae", "resume", "kişisel", "bilgiler", "iletişim"}


@dataclass
class Aday:
    dosya: str
    ad_soyad: str = ""
    eposta: str = ""
    telefon: str = ""
    linkedin: str = ""
    deneyim_yil: float = 0.0
    egitim: str = ""
    diller: list[str] = field(default_factory=list)
    beceriler: list[str] = field(default_factory=list)
    ilan_uyum: float | None = None
    uyarilar: list[str] = field(default_factory=list)


def ad_bul(metin: str, dosya_adi: str) -> str:
    for satir in metin.splitlines()[:8]:
        s = satir.strip()
        if not s or "@" in s or any(ch.isdigit() for ch in s):
            continue
        kelimeler = s.split()
        if 2 <= len(kelimeler) <= 4 and all(re.fullmatch(r"[A-Za-zÇĞİÖŞÜçğıöşüÂâÎîÛû.'-]+", k) for k in kelimeler):
            if not (set(map(kucuk, kelimeler)) & BASLIK_KELIMELERI):
                return " ".join(k[0] + kucuk(k[1:]) if k.isupper() else k for k in kelimeler)
    return Path(dosya_adi).stem.replace("_", " ").replace("-", " ").title()


def telefon_bul(metin: str) -> str:
    m = TELEFON.search(metin)
    return f"+90 {m.group(1)} {m.group(2)} {m.group(3)} {m.group(4)}" if m else ""


def deneyim_hesapla(metin: str, bugun: date | None = None) -> float:
    """Tarih aralıklarını bulur, çakışanları birleştirir ve toplam yılı döndürür."""
    bu_yil = (bugun or date.today()).year
    araliklar = []
    for m in ARALIK.finditer(metin):
        bas = int(m.group(1))
        bit = int(m.group(2)) if m.group(2) else bu_yil
        if bas <= bit <= bu_yil:
            araliklar.append((bas, bit))
    if not araliklar:
        return 0.0
    araliklar.sort()
    toplam, (cbas, cbit) = 0, araliklar[0]
    for bas, bit in araliklar[1:]:
        if bas <= cbit:
            cbit = max(cbit, bit)
        else:
            toplam += cbit - cbas
            cbas, cbit = bas, bit
    toplam += cbit - cbas
    return float(toplam)


def egitim_bul(metin_sade: str) -> str:
    for seviye, _, ifadeler in EGITIM_SEVIYELERI:
        if any(kelime_var(metin_sade, i) for i in ifadeler):
            return seviye
    return ""


def egitim_bolumunu_ayikla(metin: str) -> str:
    """Deneyim hesabında okul yıllarını saymamak için 'Eğitim' bölümünü çıkarır."""
    satirlar, cikti, egitimde = metin.splitlines(), [], False
    for satir in satirlar:
        b = sade(satir.strip()).rstrip(":")
        if b in {"egitim", "egitim bilgileri", "education", "ogrenim"}:
            egitimde = True
            continue
        if b in {"deneyim", "is deneyimi", "experience", "work experience", "beceriler", "skills",
                 "yabanci dil", "diller", "languages", "sertifikalar", "referanslar", "projeler"}:
            egitimde = False
        if not egitimde:
            cikti.append(satir)
    return "\n".join(cikti)


def beceri_listesi_oku(yol: Path) -> list[str]:
    return [s.strip() for s in txt_oku(yol).splitlines() if s.strip() and not s.lstrip().startswith("#")]


def cv_analiz(yol: Path, beceri_listesi: list[str], ilan_becerileri: list[str] | None = None) -> Aday:
    aday = Aday(dosya=yol.name)
    try:
        metin = metin_oku(yol)
    except Exception as hata:  # bozuk / şifreli dosya
        aday.uyarilar.append(f"Okunamadı: {hata.__class__.__name__}")
        return aday
    if len(metin.strip()) < 40:
        aday.uyarilar.append("Metin çıkarılamadı (taranmış görüntü PDF olabilir, OCR gerekir)")
        aday.ad_soyad = ad_bul("", yol.name)
        return aday

    ms = sade(metin)
    aday.ad_soyad = ad_bul(metin, yol.name)
    e = EPOSTA.search(metin)
    aday.eposta = e.group(0) if e else ""
    aday.telefon = telefon_bul(metin)
    li = LINKEDIN.search(metin)
    aday.linkedin = li.group(0) if li else ""
    aday.deneyim_yil = deneyim_hesapla(egitim_bolumunu_ayikla(metin))
    aday.egitim = egitim_bul(ms)
    aday.diller = [d for d, ifadeler in DILLER.items() if any(kelime_var(ms, i) for i in ifadeler)]
    aday.beceriler = [b for b in beceri_listesi if kelime_var(ms, b)]

    if ilan_becerileri:
        eslesen = [b for b in ilan_becerileri if b in aday.beceriler]
        aday.ilan_uyum = round(100 * len(eslesen) / len(ilan_becerileri), 1)

    if not aday.eposta and not aday.telefon:
        aday.uyarilar.append("İletişim bilgisi bulunamadı")
    if aday.deneyim_yil == 0:
        aday.uyarilar.append("Deneyim tarihi bulunamadı")
    return aday


# ----------------------------------------------------------------------------
# Excel çıktısı
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")


def _tablo_yaz(ws, basliklar: list[str], satirlar: list[list], genislikler: list[int]) -> None:
    ws.append(basliklar)
    for hucre in ws[1]:
        hucre.fill, hucre.font = BASLIK_DOLGU, BASLIK_YAZI
        hucre.alignment = Alignment(vertical="center")
    for satir in satirlar:
        ws.append(satir)
    for i, g in enumerate(genislikler, start=1):
        ws.column_dimensions[get_column_letter(i)].width = g
    ws.freeze_panes = "A2"
    if satirlar:
        ws.auto_filter.ref = ws.dimensions


def rapor_yaz(adaylar: list[Aday], cikti: Path, ilan_var: bool) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Adaylar"
    basliklar = ["Dosya", "Ad Soyad", "E-posta", "Telefon", "LinkedIn", "Toplam Deneyim (yıl)",
                 "En Yüksek Eğitim", "Yabancı Diller", "Beceriler", "Beceri Sayısı"]
    genislikler = [24, 22, 28, 18, 30, 12, 16, 22, 50, 10]
    if ilan_var:
        basliklar.append("İlan Uyum %")
        genislikler.append(12)
    basliklar.append("Uyarılar")
    genislikler.append(40)

    sirali = sorted(adaylar, key=lambda a: (-(a.ilan_uyum or 0), -len(a.beceriler), -a.deneyim_yil))
    satirlar = []
    for a in sirali:
        s = [a.dosya, a.ad_soyad, a.eposta, a.telefon, a.linkedin, a.deneyim_yil, a.egitim,
             ", ".join(a.diller), ", ".join(a.beceriler), len(a.beceriler)]
        if ilan_var:
            s.append(a.ilan_uyum)
        s.append("; ".join(a.uyarilar))
        satirlar.append(s)
    _tablo_yaz(ws, basliklar, satirlar, genislikler)

    oz = wb.create_sheet("Özet")
    egitim_dagilimi = {}
    for a in adaylar:
        egitim_dagilimi[a.egitim or "Belirsiz"] = egitim_dagilimi.get(a.egitim or "Belirsiz", 0) + 1
    okunan = [a for a in adaylar if not any(u.startswith(("Okunamadı", "Metin çıkarılamadı")) for u in a.uyarilar)]
    ozet = [
        ["Toplam CV", len(adaylar)],
        ["Başarıyla okunan", len(okunan)],
        ["Ortalama deneyim (yıl)", round(sum(a.deneyim_yil for a in okunan) / len(okunan), 1) if okunan else 0],
        ["Uyarı içeren CV", sum(1 for a in adaylar if a.uyarilar)],
    ]
    ozet += [[f"Eğitim: {k}", v] for k, v in sorted(egitim_dagilimi.items())]
    if ilan_var:
        ozet.append(["İlan uyumu ≥ %60 olan aday", sum(1 for a in adaylar if (a.ilan_uyum or 0) >= 60)])
    _tablo_yaz(oz, ["Gösterge", "Değer"], ozet, [32, 14])

    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


# ----------------------------------------------------------------------------
# Komut satırı
# ----------------------------------------------------------------------------

def calistir(girdi: Path, cikti: Path, beceriler: Path, ilan: Path | None = None) -> list[Aday]:
    if not girdi.is_dir():
        raise SystemExit(f"Girdi klasörü bulunamadı: {girdi}")
    dosyalar = sorted(p for p in girdi.iterdir() if p.suffix.lower() in DESTEKLENEN)
    if not dosyalar:
        raise SystemExit(f"{girdi} içinde PDF, DOCX veya TXT dosyası yok.")

    beceri_listesi = beceri_listesi_oku(beceriler)
    ilan_becerileri = None
    if ilan:
        ilan_sade = sade(txt_oku(ilan))
        ilan_becerileri = [b for b in beceri_listesi if kelime_var(ilan_sade, b)]
        if not ilan_becerileri:
            print("[!] İlanda beceri listesindeki hiçbir ifade bulunamadı; uyum yüzdesi hesaplanmayacak.")
            ilan_becerileri = None

    adaylar = [cv_analiz(p, beceri_listesi, ilan_becerileri) for p in dosyalar]
    rapor_yaz(adaylar, cikti, ilan_var=ilan_becerileri is not None)
    return adaylar


def main(argv: list[str] | None = None) -> None:
    # Windows konsolu (cp1254) desteklemediği karakterde çökmesin
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="CV'leri okuyup tek bir Excel raporunda toplar.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "cvler", help="CV klasörü (PDF/DOCX/TXT)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "aday_raporu.xlsx", help="Excel çıktı yolu")
    p.add_argument("--beceriler", type=Path, default=BURASI / "beceriler.txt", help="Aranacak beceriler (her satıra bir tane)")
    p.add_argument("--ilan", type=Path, default=None, help="İsteğe bağlı iş ilanı metni (.txt)")
    a = p.parse_args(argv)

    adaylar = calistir(a.girdi, a.cikti, a.beceriler, a.ilan)
    uyarili = sum(1 for x in adaylar if x.uyarilar)
    print(f"[OK] {len(adaylar)} CV işlendi ({uyarili} tanesinde uyarı var)")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
