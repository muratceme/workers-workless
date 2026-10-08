import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))

import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"
KURLAR = {"USD": D("41.20"), "EUR": D("48.05")}


def calistir(tmp, **kw):
    return main.calistir(ORNEK / "acik_siparisler.csv", Path(tmp) / "s.xlsx", ORNEK / "teslim_gecmisi.csv", KURLAR, date(2026, 10, 8), **kw)


class OrnekTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        with tempfile.TemporaryDirectory() as tmp:
            cls.s = calistir(tmp)
            cls.wb = load_workbook(Path(tmp) / "s.xlsx")
        cls.d = {x.ref: x for x in cls.s["satirlar"]}

    def test_gecikenler(self):
        gec = {r: (x.oncelik, x.gecikme) for r, x in self.d.items() if x.durum == "Gecikmiş"}
        self.assertEqual(gec, {"SA-26-0090/10": ("Kritik", 37), "SA-26-0101/10": ("Yüksek", 13), "SA-26-0101/20": ("Yüksek", 8),
                               "SA-26-0105/10": ("Gecikmiş", 7), "SA-26-0110/10": ("Yüksek", 23), "SA-26-0110/20": ("Yüksek", 23),
                               "SA-26-0115/10": ("Gecikmiş", 3)})
        # 0101/20: istenen termin 25.09, teyit 30.09 → gecikme teyitten sayılır (8 gün)
        self.assertIn("Teyit termini istenenden 5 gün sonra", self.d["SA-26-0101/20"].notlar)

    def test_tutarlar(self):
        self.assertEqual(self.d["SA-26-0090/10"].kalan_tl, D("376712.00"))        # 800 kg × 9,80 EUR × 48,05
        self.assertEqual(self.d["SA-26-0105/10"].kalan_tl, D("93936.00"))         # kalan 60 × 38 USD × 41,20
        self.assertEqual(self.d["SA-26-0101/10"].kalan_tl, D("15600.00"))         # kalan 300 × 52
        gec = sum(x.kalan_tl for x in self.s["satirlar"] if x.durum == "Gecikmiş")
        self.assertEqual(gec, D("645848.00"))

    def test_riskler(self):
        risk = {r for r, x in self.d.items() if x.durum == "Riskli"}
        self.assertEqual(risk, {"SA-26-0125/10", "SA-26-0128/10", "SA-26-0130/10", "SA-26-0131/10"})
        # Örnek Metal geçmiş gecikmeleri 5, 12, 0, 8, 20 → ort. 9 gün → 12.10 + 9 = 21.10 > ihtiyaç 15.10
        self.assertEqual(self.d["SA-26-0130/10"].tahmini, date(2026, 10, 21))
        self.assertTrue(self.d["SA-26-0130/10"].notlar[0].startswith("Tahmini teslim (21.10.2026) ihtiyaç tarihinden"))
        self.assertTrue(self.d["SA-26-0125/10"].notlar[0].startswith("Termin ihtiyaç tarihinden (20.10.2026) 3 gün sonra"))
        self.assertIn("Termin yaklaşıyor, teyit yok", self.d["SA-26-0131/10"].notlar)
        self.assertEqual(self.d["SA-26-0120/10"].durum, "Yaklaşan")                 # Kurgu Plastik zamanında teslim ediyor
        self.assertEqual(self.d["SA-26-0135/10"].notlar, [])                        # 3 gündür teyitsiz: henüz uyarı yok

    def test_gecmis(self):
        g = main.gecmis_oku(ORNEK / "teslim_gecmisi.csv")
        p = g[main.katla("Kurgu Plastik A.Ş.")]
        self.assertEqual((p["adet"], p["zamaninda"], p["ort"]), (5, 0.8, D("0.2")))

    def test_kontroller_ve_ozet(self):
        k = {(o, r) for o, r, _ in self.s["kontroller"]}
        self.assertIn(("Hata", "SA-26-0136/10"), k)
        self.assertIn(("Bilgi", "SA-26-0118/10"), k)                               # %5 fazla teslim
        self.assertEqual([t["tedarikci"] for t in self.s["ozet"]][:2], ["Kurgu Kimya Ltd. Şti.", "Örnek Ambalaj A.Ş."])
        self.assertEqual(set(self.s["hatirlatmalar"]), {"Kurgu Kimya Ltd. Şti.", "Örnek Ambalaj A.Ş.", "Deneme Elektrik Ltd. Şti.",
                                                         "Örnek Metal San. Ltd. Şti."})
        self.assertIn("en geç 10.10.2026", self.s["hatirlatmalar"]["Örnek Metal San. Ltd. Şti."])

    def test_excel(self):
        self.assertEqual(self.wb.sheetnames, ["Özet", "Gecikenler", "Gecikme Riski", "Termin Takvimi", "Açık Siparişler", "Tedarikçi Hatırlatma",
                                              "Kontroller"])
        g = self.wb["Gecikenler"]
        self.assertEqual([g.cell(i, 2).value for i in range(2, 5)], ["SA-26-0090/10", "SA-26-0110/10", "SA-26-0110/20"])
        t = self.wb["Termin Takvimi"]
        toplam = [r for r in t.iter_rows(values_only=True) if r[0] == "Toplam"][0]
        self.assertEqual(toplam[1], 645848.0)                                      # geçmiş termin = geciken tutar


class OkumaTesti(unittest.TestCase):
    def test_kur_eksik_ve_baslik(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "s.csv"
            y.write_text("Açık Siparişler Raporu\n\nPO No,Firma,Miktar,Gelen Miktar,Termin Tarihi,Fiyat,Döviz\n"
                         "P1,A,10,0,01.10.2026,5,GBP\nP2,B,1.250,0,20.10.2026,2,TL\n", encoding="utf-8")
            s = main.calistir(y, Path(tmp) / "s.xlsx", bugun=date(2026, 10, 8))
        p1, p2 = s["satirlar"]
        self.assertEqual((p1.durum, p1.kalan_tl), ("Gecikmiş", None))
        self.assertEqual((p2.miktar, p2.kalan_tl, p2.durum), (D(1250), D("2500.00"), "Riskli"))                  # 12 gün kaldı, teyit yok
        self.assertIn(("Dikkat", "GBP"), {(o, r) for o, r, _ in s["kontroller"]})

    def test_para(self):
        self.assertEqual(main.para("1.250,50"), D("1250.50"))
        self.assertEqual(main.para("9,80"), D("9.80"))
        self.assertEqual(main.para("0.95"), D("0.95"))


if __name__ == "__main__":
    unittest.main()
