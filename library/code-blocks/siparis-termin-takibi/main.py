"""
Sipariş Termin Takibi — Workers / Workless kod bloğu
Satın Alma › Satın Alma Uzman Yardımcısı

Açık satın alma siparişlerini termin tarihine göre izler:
  - Geçerli termin: tedarikçi teyit termini varsa o, yoksa sipariş termini.
  - Geciken satırlar: kalan miktarı olan ve geçerli termini geçmiş satırlar; gecikme gününe göre öncelik
    (1-7 gün Gecikmiş, 8-30 gün Yüksek, 30+ gün Kritik) ve kalan tutar (dövizliler --kur ile TL).
  - Gecikme riski: önümüzdeki --ufuk günü içinde termini gelen satırlarda teyit yoksa, tedarikçinin geçmiş
    teslimlerindeki ortalama gecikmeye göre tahmini teslim termini aşıyorsa veya ihtiyaç tarihini geçiyorsa.
  - Teyit termini istenen terminden sonra olan, --teyit-gun günden uzun süredir teyitsiz bekleyen, fazla teslim
    alınan ve termini boş satırlar.
  - Tedarikçi özeti (açık/geciken tutar, geçmiş zamanında teslim oranı), haftalık termin takvimi ve tedarikçiye
    gönderilebilecek hatırlatma metni taslağı. Hiçbir şey gönderilmez.
İnternete bağlanmaz.

Kullanım:
    python main.py                                                    # örnek veriyle (rapor tarihi 08.10.2026)
    python main.py --girdi acik_siparisler.xlsx --gecmis teslimler.xlsx --kur USD=41,20 EUR=48,05
    python main.py --girdi acik_siparisler.xlsx --tarih 15.10.2026 --ufuk 10 --teyit-gun 3
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri" / "acik_siparisler.csv"
SIFIR = Decimal(0)
KURUS = Decimal("0.01")
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
TL_ADLARI = {"", "tl", "try", "tl.", "₺"}

SUTUNLAR = {
    "no": ("siparis no", "siparis numarasi", "po no", "belge no", "siparis"),
    "satir": ("satir", "satir no", "kalem", "kalem no", "pozisyon"),
    "tedarikci": ("tedarikci", "tedarikci adi", "firma", "satici"),
    "kod": ("malzeme kodu", "stok kodu", "urun kodu", "kod"),
    "malzeme": ("malzeme", "malzeme adi", "malzeme aciklamasi", "urun", "aciklama"),
    "siparis_tarihi": ("siparis tarihi", "belge tarihi", "tarih"),
    "termin": ("termin", "termin tarihi", "istenen termin", "teslim tarihi", "istenen teslim tarihi"),
    "teyit": ("teyit termini", "teyit tarihi", "onaylanan termin", "tedarikci termini", "teyitli termin"),
    "miktar": ("siparis miktari", "miktar"),
    "teslim": ("teslim alinan", "teslim alinan miktar", "gelen miktar", "teslim miktari", "giris miktari"),
    "birim": ("birim", "olcu birimi"),
    "fiyat": ("birim fiyat", "fiyat", "net fiyat"),
    "doviz": ("para birimi", "doviz", "pb", "doviz cinsi"),
    "ihtiyac": ("ihtiyac tarihi", "uretim tarihi", "kullanim tarihi"),
}
GECMIS_SUTUNLAR = {"tedarikci": SUTUNLAR["tedarikci"], "termin": ("termin", "termin tarihi", "teyit termini"),
                   "teslim": ("teslim tarihi", "giris tarihi", "irsaliye tarihi", "gelis tarihi")}


def katla(s) -> str:
    return " ".join(str(s or "").translate(_TR).lower().split())


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("₺", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1 or (s.count(".") == 1 and len(s.split(".")[1]) == 3 and s.split(".")[0] not in ("", "0")):
        s = s.replace(".", "")                                  # 1.250 → 1250 (binlik)
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%y", "%Y-%m-%d %H:%M:%S", "%d.%m.%Y %H:%M"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def tl(x: Decimal) -> str:
    return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def basliklari_bul(s: list[list], sutunlar: dict, zorunlu: tuple, ad: str) -> tuple[int, dict]:
    for bi, r in enumerate(s[:10]):                            # üstteki rapor başlıklarını atla
        b = [katla(x) for x in r]
        k = {a: next((i for i, x in enumerate(b) if x in es), None) for a, es in sutunlar.items()}
        if all(k[z] is not None for z in zorunlu):
            return bi, k
    raise SystemExit(f"{ad}: gerekli sütunlar bulunamadı ({', '.join(zorunlu)}). İlk satır: {s[0] if s else '(boş)'}")


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Satir:
    no: str
    satir: str
    tedarikci: str
    kod: str
    malzeme: str
    siparis_tarihi: date | None
    termin: date | None
    teyit: date | None
    miktar: Decimal
    teslim: Decimal
    birim: str
    fiyat: Decimal | None
    doviz: str
    ihtiyac: date | None
    kalan_tl: Decimal | None = None
    durum: str = ""
    oncelik: str = ""
    gecikme: int = 0
    tahmini: date | None = None
    notlar: list[str] = field(default_factory=list)

    @property
    def ref(self) -> str:
        return f"{self.no}/{self.satir}" if self.satir else self.no

    @property
    def gecerli(self) -> date | None:
        return self.teyit or self.termin

    @property
    def kalan(self) -> Decimal:
        return max(self.miktar - self.teslim, SIFIR)


def siparisleri_oku(yol: Path) -> tuple[list[Satir], list[tuple[str, str, str]]]:
    s = tablo_oku(yol)
    bi, k = basliklari_bul(s, SUTUNLAR, ("no", "tedarikci", "miktar"), yol.name)
    al = lambda r, a: r[k[a]] if k[a] is not None and k[a] < len(r) else None  # noqa: E731
    satirlar, kontroller = [], []
    for r in s[bi + 1:]:
        no = str(al(r, "no") or "").strip()
        if not no:
            continue
        miktar = para(al(r, "miktar"))
        x = Satir(no, str(al(r, "satir") or "").strip(), str(al(r, "tedarikci") or "").strip() or "(tedarikçi yok)",
                  str(al(r, "kod") or "").strip(), str(al(r, "malzeme") or "").strip(), tarih(al(r, "siparis_tarihi")),
                  tarih(al(r, "termin")), tarih(al(r, "teyit")), miktar or SIFIR, para(al(r, "teslim")) or SIFIR,
                  str(al(r, "birim") or "").strip(), para(al(r, "fiyat")), str(al(r, "doviz") or "").strip().upper() or "TL",
                  tarih(al(r, "ihtiyac")))
        if miktar is None or miktar <= 0:
            kontroller.append(("Hata", x.ref, f"Sipariş miktarı okunamadı veya sıfır ({al(r, 'miktar')!r})"))
            continue
        if x.gecerli is None:
            kontroller.append(("Hata", x.ref, "Termin tarihi boş veya okunamadı; takibe alınamadı"))
        if x.fiyat is None:
            kontroller.append(("Dikkat", x.ref, "Birim fiyat yok; tutar hesaplanamadı"))
        if x.siparis_tarihi and x.termin and x.termin < x.siparis_tarihi:
            kontroller.append(("Hata", x.ref, f"Termin ({x.termin:%d.%m.%Y}) sipariş tarihinden önce"))
        satirlar.append(x)
    return satirlar, kontroller


def gecmis_oku(yol: Path) -> dict[str, dict]:
    """Tedarikçi bazında geçmiş teslim performansı: teslim sayısı, zamanında oranı, ortalama gecikme (erken = 0)."""
    s = tablo_oku(yol)
    bi, k = basliklari_bul(s, GECMIS_SUTUNLAR, ("tedarikci", "termin", "teslim"), yol.name)
    gun = defaultdict(list)
    for r in s[bi + 1:]:
        ted = str(r[k["tedarikci"]] if k["tedarikci"] < len(r) else "").strip()
        t, g = tarih(r[k["termin"]]), tarih(r[k["teslim"]])
        if ted and t and g:
            gun[katla(ted)].append((g - t).days)
    return {ted: {"adet": len(g), "zamaninda": sum(1 for x in g if x <= 0) / len(g),
                  "ort": Decimal(sum(max(x, 0) for x in g)) / len(g)} for ted, g in gun.items()}


# ----------------------------------------------------------------------------
# Hesap
# ----------------------------------------------------------------------------

def oncelik(gun: int) -> str:
    return "Kritik" if gun > 30 else "Yüksek" if gun > 7 else "Gecikmiş"


def degerlendir(satirlar: list[Satir], gecmis: dict, kurlar: dict[str, Decimal], bugun: date, ufuk: int, teyit_gun: int,
                kontroller: list) -> None:
    eksik_kur = set()
    for x in satirlar:
        if x.fiyat is not None:
            kur = Decimal(1) if katla(x.doviz) in TL_ADLARI else kurlar.get(x.doviz)
            if kur is None:
                eksik_kur.add(x.doviz)
            else:
                x.kalan_tl = (x.kalan * x.fiyat * kur).quantize(KURUS, ROUND_HALF_UP)
        if x.teslim > x.miktar:
            x.notlar.append(f"Fazla teslim: {x.teslim - x.miktar:g} {x.birim} (%{(x.teslim / x.miktar - 1) * 100:.1f})".replace(".", ","))
        if x.kalan <= 0:
            x.durum = "Tamamlandı"
            continue
        if x.gecerli is None:
            x.durum = "Termin yok"
            continue
        p = gecmis.get(katla(x.tedarikci))
        ort = p["ort"] if p and p["adet"] >= 3 else SIFIR
        x.tahmini = x.gecerli + timedelta(days=int(ort.quantize(Decimal(1), ROUND_HALF_UP)))
        if x.teyit and x.termin and x.teyit > x.termin:
            x.notlar.append(f"Teyit termini istenenden {(x.teyit - x.termin).days} gün sonra")
        if not x.teyit and x.siparis_tarihi and (bugun - x.siparis_tarihi).days > teyit_gun:
            x.notlar.append(f"Tedarikçi teyidi yok ({(bugun - x.siparis_tarihi).days} gündür)")
        if x.teslim > 0:
            x.notlar.append(f"Kısmi teslim: {x.teslim:g}/{x.miktar:g} {x.birim}")
        if x.gecerli < bugun:
            x.durum, x.gecikme = "Gecikmiş", (bugun - x.gecerli).days
            x.oncelik, x.tahmini = oncelik(x.gecikme), None
            if x.ihtiyac and x.ihtiyac < bugun:
                x.notlar.append(f"İhtiyaç tarihi ({x.ihtiyac:%d.%m.%Y}) geçti")
            continue
        riskler = []
        if x.ihtiyac and x.gecerli > x.ihtiyac:
            riskler.append(f"Termin ihtiyaç tarihinden ({x.ihtiyac:%d.%m.%Y}) {(x.gecerli - x.ihtiyac).days} gün sonra")
        elif x.ihtiyac and x.tahmini > x.ihtiyac:
            riskler.append(f"Tahmini teslim ({x.tahmini:%d.%m.%Y}) ihtiyaç tarihinden ({x.ihtiyac:%d.%m.%Y}) sonra")
        if (x.gecerli - bugun).days <= ufuk:
            if not x.teyit:
                riskler.append("Termin yaklaşıyor, teyit yok")
            if (x.tahmini - x.gecerli).days >= 3:
                riskler.append(f"Tedarikçi geçmişte ortalama {str(ort.quantize(Decimal('0.1'))).replace('.', ',')} gün gecikmeli teslim etmiş "
                               f"(tahmini {x.tahmini:%d.%m.%Y})")
        if riskler:
            x.durum, x.oncelik = "Riskli", "Riskli"
            x.notlar = riskler + x.notlar
        else:
            x.durum = "Yaklaşan" if (x.gecerli - bugun).days <= ufuk else "Zamanında"
    for d in sorted(eksik_kur):
        kontroller.append(("Dikkat", d, f"{d} kuru verilmedi (--kur {d}=...); bu para birimindeki satırların TL tutarı boş"))


def tedarikci_ozeti(satirlar: list[Satir], gecmis: dict) -> list[dict]:
    t = defaultdict(lambda: {"acik": 0, "acik_tl": SIFIR, "gec": 0, "gec_tl": SIFIR, "risk": 0, "en_uzun": 0, "teyitsiz": 0})
    for x in satirlar:
        if x.durum in ("Tamamlandı", "Termin yok"):
            continue
        o = t[x.tedarikci]
        o["acik"] += 1
        o["acik_tl"] += x.kalan_tl or SIFIR
        if x.durum == "Gecikmiş":
            o["gec"] += 1
            o["gec_tl"] += x.kalan_tl or SIFIR
            o["en_uzun"] = max(o["en_uzun"], x.gecikme)
        o["risk"] += x.durum == "Riskli"
        o["teyitsiz"] += any(n.startswith("Tedarikçi teyidi yok") for n in x.notlar)
    sonuc = []
    for ted, o in t.items():
        p = gecmis.get(katla(ted))
        sonuc.append({"tedarikci": ted, **o, "gecmis": p})
    return sorted(sonuc, key=lambda o: (-o["gec_tl"], -o["gec"], o["tedarikci"]))


def hatirlatma_metni(ted: str, satirlar: list[Satir], cevap_tarihi: date) -> str:
    gec = [x for x in satirlar if x.tedarikci == ted and x.durum == "Gecikmiş"]
    teyit = [x for x in satirlar if x.tedarikci == ted and x.durum != "Gecikmiş" and any(n.startswith("Tedarikçi teyidi yok") for n in x.notlar)]
    if not gec and not teyit:
        return ""
    sat = [f"Konu: Açık siparişlerimiz hk. — {ted}", "", "Sayın Yetkili,", ""]
    if gec:
        sat.append("Aşağıdaki sipariş satırlarımızın teslim tarihi geçmiştir:")
        sat += [f"- {x.ref} · {x.kod} {x.malzeme} · kalan {x.kalan:g} {x.birim} · termin {x.gecerli:%d.%m.%Y} ({x.gecikme} gün)" for x in gec]
        sat.append("")
    if teyit:
        sat.append("Aşağıdaki siparişlerimiz için termin teyidinizi henüz alamadık:")
        sat += [f"- {x.ref} · {x.kod} {x.malzeme} · {x.kalan:g} {x.birim} · istenen termin {x.termin:%d.%m.%Y}" for x in teyit if x.termin]
        sat.append("")
    sat += [f"Güncel sevk tarihlerini en geç {cevap_tarihi:%d.%m.%Y} tarihine kadar bildirmenizi rica ederiz.", "", "Saygılarımızla,",
            "[Ad Soyad] · Satın Alma"]
    return "\n".join(sat)


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Kritik": "F8C9C6", "Yüksek": "FDE2E1", "Gecikmiş": "FFF4CE", "Riskli": "FFE8CC", "Hata": "FDE2E1", "Dikkat": "FFF4CE", "Bilgi": "E8F0FE"}
UST = Alignment(vertical="top", wrap_text=True)
KALIN = Font(bold=True)
SIRA = {"Kritik": 0, "Yüksek": 1, "Gecikmiş": 2, "Riskli": 3}


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def _tarih_bicim(ws, sutunlar):
    for c in sutunlar:
        ws.cell(ws.max_row, c).number_format = "DD.MM.YYYY"


SATIR_BASLIK = ["Öncelik", "Sipariş/Satır", "Tedarikçi", "Malzeme Kodu", "Malzeme", "Sipariş Tarihi", "Termin", "Teyit Termini", "Gecikme (gün)",
                "Kalan", "Birim", "Kalan Tutar (TL)", "Tahmini Teslim", "İhtiyaç Tarihi", "Notlar", "Aksiyon / Yeni Termin"]
SATIR_GENISLIK = (10, 15, 26, 11, 24, 11, 11, 11, 9, 9, 7, 15, 12, 12, 50, 22)


def _satir_yaz(ws, x: Satir):
    ws.append([x.oncelik or x.durum, x.ref, x.tedarikci, x.kod, x.malzeme, x.siparis_tarihi, x.termin, x.teyit, x.gecikme or None,
               float(x.kalan), x.birim, None if x.kalan_tl is None else float(x.kalan_tl), x.tahmini, x.ihtiyac, "\n".join(x.notlar), ""])
    _tarih_bicim(ws, (6, 7, 8, 13, 14))
    ws.cell(ws.max_row, 12).number_format = "#,##0.00"
    if (x.oncelik or x.durum) in RENK:
        ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=RENK[x.oncelik or x.durum])
    ws.cell(ws.max_row, 16).fill = PatternFill("solid", fgColor="FFF4CE")
    for h in ws[ws.max_row]:
        h.alignment = UST


def rapor_yaz(cikti: Path, satirlar: list[Satir], ozet: list[dict], kontroller: list, bugun: date, ufuk: int, hatirlatmalar: dict) -> None:
    wb = Workbook()
    o = wb.active
    o.title = "Özet"
    acik = [x for x in satirlar if x.durum not in ("Tamamlandı", "Termin yok")]
    gec = [x for x in acik if x.durum == "Gecikmiş"]
    toplam = lambda xs: sum((x.kalan_tl or SIFIR for x in xs), SIFIR)  # noqa: E731
    o.column_dimensions["A"].width = 34
    for etiket, deger in [("Rapor tarihi", bugun), ("Açık sipariş satırı", len(acik)), ("Açık tutar (TL)", float(toplam(acik))),
                          ("Geciken satır", len(gec)), ("Geciken tutar (TL)", float(toplam(gec))),
                          ("Kritik (30+ gün)", sum(1 for x in gec if x.oncelik == "Kritik")),
                          (f"Gecikme riski ({ufuk} gün içinde / ihtiyaç tarihi)", sum(1 for x in acik if x.durum == "Riskli"))]:
        o.append([etiket, deger])
        o.cell(o.max_row, 1).font = KALIN
        o.cell(o.max_row, 2).number_format = "DD.MM.YYYY" if isinstance(deger, date) else "#,##0.00" if isinstance(deger, float) else "General"
    o.append([])
    _baslik(o, ["Tedarikçi", "Açık Satır", "Açık Tutar (TL)", "Geciken Satır", "Geciken Tutar (TL)", "En Uzun Gecikme (gün)", "Riskli",
                "Teyitsiz", "Geçmiş Teslim", "Zamanında Oranı", "Ort. Gecikme (gün)"], (34, 10, 15, 12, 16, 12, 8, 9, 12, 13, 13))
    for t in ozet:
        p = t["gecmis"]
        o.append([t["tedarikci"], t["acik"], float(t["acik_tl"]), t["gec"], float(t["gec_tl"]), t["en_uzun"] or None, t["risk"], t["teyitsiz"],
                  p["adet"] if p else "geçmiş yok", p["zamaninda"] if p else None, float(p["ort"].quantize(Decimal("0.1"))) if p else None])
        for c in (3, 5):
            o.cell(o.max_row, c).number_format = "#,##0.00"
        o.cell(o.max_row, 10).number_format = "0%"
        if t["gec"]:
            o.cell(o.max_row, 5).fill = PatternFill("solid", fgColor=RENK["Yüksek"])
    o.append([])
    o.append(["Not", "Tutarlar kalan miktar × birim fiyat × kur ile hesaplanır. Ortalama gecikme geçmiş teslimlerden (erken teslim 0 sayılır) "
                     "ve en az 3 teslim varsa tahmini teslime eklenir. Öncelik eşikleri: 1-7 gün Gecikmiş, 8-30 Yüksek, 30+ Kritik."])

    g = wb.create_sheet("Gecikenler")
    _baslik(g, SATIR_BASLIK, SATIR_GENISLIK)
    for x in sorted(gec, key=lambda x: (SIRA[x.oncelik], -x.gecikme)):
        _satir_yaz(g, x)
    g.freeze_panes = "C2"
    r = wb.create_sheet("Gecikme Riski")
    _baslik(r, SATIR_BASLIK, SATIR_GENISLIK)
    for x in sorted((x for x in acik if x.durum == "Riskli"), key=lambda x: x.gecerli):
        _satir_yaz(r, x)
    r.freeze_panes = "C2"

    t = wb.create_sheet("Termin Takvimi")
    hafta = lambda d: d - timedelta(days=d.weekday())  # noqa: E731
    haftalar = [hafta(bugun) + timedelta(days=7 * i) for i in range(6)]
    teds = sorted({x.tedarikci for x in acik})
    _baslik(t, ["Tedarikçi \\ Hafta (bekleyen tutar, TL)", "Termini geçmiş"] + [f"{h:%d.%m} haftası" for h in haftalar] + ["Sonra"], (34, 14) + (13,) * 7)
    for ted in teds + ["Toplam"]:
        xs = [x for x in acik if ted in ("Toplam", x.tedarikci)]
        satir = [ted, float(toplam(x for x in xs if x.gecerli < bugun))]
        satir += [float(toplam(x for x in xs if max(h, bugun) <= x.gecerli < h + timedelta(days=7))) for h in haftalar]
        satir.append(float(toplam(x for x in xs if x.gecerli >= haftalar[-1] + timedelta(days=7))))
        t.append([satir[0]] + [v or None for v in satir[1:]])
        for c in range(2, 10):
            t.cell(t.max_row, c).number_format = "#,##0"
        if ted == "Toplam":
            for h in t[t.max_row]:
                h.font = KALIN

    a = wb.create_sheet("Açık Siparişler")
    _baslik(a, ["Durum"] + SATIR_BASLIK[1:], (11,) + SATIR_GENISLIK[1:])
    for x in sorted(satirlar, key=lambda x: (x.tedarikci, x.gecerli or date.max)):
        if x.durum != "Tamamlandı" or x.notlar:
            _satir_yaz(a, x)
            a.cell(a.max_row, 1).value = x.durum
    a.freeze_panes = "C2"
    a.auto_filter.ref = a.dimensions

    h = wb.create_sheet("Tedarikçi Hatırlatma")
    _baslik(h, ["Tedarikçi", "Hatırlatma metni (taslak — gönderilmez)", "Gönderildi mi?"], (30, 100, 14))
    for ted, metin in hatirlatmalar.items():
        h.append([ted, metin, ""])
        h.cell(h.max_row, 3).fill = PatternFill("solid", fgColor="FFF4CE")
        for c in h[h.max_row]:
            c.alignment = UST

    k = wb.create_sheet("Kontroller")
    _baslik(k, ["Önem", "Satır / Konu", "Açıklama"], (9, 18, 90))
    for x in satirlar:
        for n in x.notlar:
            if n.startswith("Fazla teslim") and x.durum == "Tamamlandı":
                kontroller.append(("Bilgi", x.ref, n))
    for kt in kontroller:
        k.append(list(kt))
        k.cell(k.max_row, 1).fill = PatternFill("solid", fgColor=RENK[kt[0]])
    if not kontroller:
        k.append(["Bilgi", "-", "Veri kontrollerinde sorun bulunmadı"])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


# ----------------------------------------------------------------------------
# Akış
# ----------------------------------------------------------------------------

def calistir(girdi: Path, cikti: Path, gecmis_yolu: Path | None = None, kurlar: dict[str, Decimal] | None = None, bugun: date | None = None,
             ufuk: int = 14, teyit_gun: int = 5, cevap_gun: int = 2) -> dict:
    bugun = bugun or date.today()
    satirlar, kontroller = siparisleri_oku(girdi)
    gecmis = gecmis_oku(gecmis_yolu) if gecmis_yolu else {}
    degerlendir(satirlar, gecmis, kurlar or {}, bugun, ufuk, teyit_gun, kontroller)
    ozet = tedarikci_ozeti(satirlar, gecmis)
    hatirlatmalar = {t["tedarikci"]: m for t in ozet if (m := hatirlatma_metni(t["tedarikci"], satirlar, bugun + timedelta(days=cevap_gun)))}
    rapor_yaz(cikti, satirlar, ozet, kontroller, bugun, ufuk, hatirlatmalar)
    return {"satirlar": satirlar, "ozet": ozet, "kontroller": kontroller, "hatirlatmalar": hatirlatmalar, "bugun": bugun}


def kur_coz(degerler: list[str]) -> dict[str, Decimal]:
    kurlar = {}
    for x in degerler:
        ad, _, v = x.partition("=")
        if para(v) is None:
            raise SystemExit(f"--kur 'USD=41,20' biçiminde olmalı: {x}")
        kurlar[ad.strip().upper()] = para(v)
    return kurlar


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Açık satın alma siparişlerini termine göre izler; geciken ve riskli satırları tedarikçi bazında listeler.")
    p.add_argument("--girdi", type=Path, default=ORNEK, help="Açık siparişler (.xlsx/.csv): Sipariş No, Tedarikçi, Sipariş Miktarı, Termin...")
    p.add_argument("--gecmis", type=Path, help="Geçmiş teslimler (Tedarikçi, Termin, Teslim Tarihi) — gecikme riski için")
    p.add_argument("--kur", nargs="*", default=None, metavar="PB=KUR", help="Döviz kurları, ör. USD=41,20 EUR=48,05")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY (varsayılan: bugün)")
    p.add_argument("--ufuk", type=int, default=14, help="Gecikme riski için bakılacak gün (varsayılan 14)")
    p.add_argument("--teyit-gun", type=int, default=5, help="Siparişten sonra teyit beklenen en fazla gün (varsayılan 5)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "siparis_termin_takibi.xlsx")
    a = p.parse_args(argv)
    gecmis, kurlar, bugun = a.gecmis, kur_coz(a.kur or []), tarih(a.tarih) if a.tarih else None
    if a.tarih and bugun is None:
        print(f"[X] Tarih GG.AA.YYYY olmalı: {a.tarih}")
        return 2
    if a.girdi == ORNEK:                                        # örnek veri: sabit tarih, örnek kurlar ve geçmiş
        ayar = json.loads((BURASI / "ornek_veri" / "ayarlar.json").read_text(encoding="utf-8"))
        bugun = bugun or tarih(ayar["tarih"])
        kurlar = kurlar or {k: para(v) for k, v in ayar["kur"].items()}
        gecmis = gecmis or BURASI / "ornek_veri" / "teslim_gecmisi.csv"
        print(f"[i] Örnek veri: rapor tarihi {bugun:%d.%m.%Y}, kurlar örnektir")
    s = calistir(a.girdi, a.cikti, gecmis, kurlar, bugun, a.ufuk, a.teyit_gun)
    acik = [x for x in s["satirlar"] if x.durum not in ("Tamamlandı", "Termin yok")]
    gec = [x for x in acik if x.durum == "Gecikmiş"]
    tutar = sum((x.kalan_tl or SIFIR for x in gec), SIFIR)
    print(f"[OK] {len(acik)} açık satır · {len(gec)} geciken ({tl(tutar)} TL) · {sum(1 for x in acik if x.durum == 'Riskli')} riskli")
    for t in s["ozet"][:5]:
        if t["gec"]:
            print(f"     {t['tedarikci']}: {t['gec']} geciken satır, {tl(t['gec_tl'])} TL, en uzun {t['en_uzun']} gün")
    for o, ref, aciklama in s["kontroller"]:
        if o != "Bilgi":
            print(f"[{'X' if o == 'Hata' else '!'}] {ref}: {aciklama}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
