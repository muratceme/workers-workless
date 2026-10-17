"""
Mağazalar Arası Transfer Önerisi — Gerisi Bende kod bloğu
Ürün Planlama ve Alokasyon › Alokasyon Uzmanı

Fazla stoklu ve tükenme riskli mağazaları eşleştirip beden-renk bazında transfer önerisi çıkarır.
  Haftalık satış  = dönem satış adedi × 7 ÷ dönem gün sayısı (--donem-gun, varsayılan 28)
  Kapsama (hafta) = stok ÷ haftalık satış
  Alıcı           = satışı var ve kapsama < --min-hafta (2);  ihtiyaç = hedef hafta × haftalık satış − stok (yukarı yuvarlanır)
  Verici          = kapsama > --max-hafta (8) veya satışı yok;  fazla = stok − hedef hafta × haftalık satış (aşağı yuvarlanır), --asgari-stok adet mağazada bırakılır
  Eşleştirme      = her ürün-beden-renk için alıcılar en düşük kapsamadan başlar; önce aynı bölgedeki, sonra fazlası en çok olan vericiden alınır (sezgisel)
Nakliye maliyeti, koli içi adet ve mağaza kapasitesi hesaba katılmaz; öneri karar değildir.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"magaza": ("magaza", "magaza adi", "sube"), "bolge": ("bolge", "il", "sehir"), "urun": ("urun", "urun kodu", "model", "sku", "stok kodu"), "beden": ("beden",), "renk": ("renk",),
         "stok": ("stok", "stok adedi", "mevcut stok"), "satis": ("satis", "satis adedi", "donem satis", "donem satis adedi")}
Z = Decimal(0)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("magaza", "urun", "stok", "satis")):
        m, u = metin(r["magaza"]), metin(r["urun"])
        if not m or not u:
            continue
        kim = f"{m} · {u} (satır {r['_satir']})"
        stok, satis = para(r["stok"]), para(r["satis"])
        if stok is None or satis is None or stok < 0 or satis < 0:
            uy.append(uyari("Yüksek", "Değer okunamadı", kim, f"Stok '{metin(r['stok'])}', satış '{metin(r['satis'])}'. Satır alınmadı."))
            continue
        x = {"magaza": m, "bolge": metin(r["bolge"]), "urun": u, "beden": metin(r["beden"]), "renk": metin(r["renk"]), "stok": stok, "satis": satis}
        a = tuple(katla(x[k]) for k in ("magaza", "urun", "beden", "renk"))
        if a in gorulen:
            uy.append(uyari("Yüksek", "Mükerrer satır", kim, "Aynı mağaza, ürün, beden ve renk birden fazla satırda; yalnız ilki alındı."))
            continue
        gorulen.add(a)
        l.append(x)
    return l, uy


def hafta(x: Decimal | None) -> str:
    return "satış yok" if x is None else f"{x:f}"


def oner(satirlar: list[dict], donem_gun: Decimal = Decimal(28), hedef: Decimal = Decimal(4), alt: Decimal = Decimal(2), ust: Decimal = Decimal(8), asgari: Decimal = Z, min_transfer: Decimal = Decimal(1)) -> dict:
    if not satirlar:
        raise ValueError("Değerlendirilecek stok satırı yok.")
    if not alt <= hedef <= ust:
        raise ValueError("Hedef hafta, asgari ve azami kapsama haftalarının arasında olmalı (min ≤ hedef ≤ max).")
    sku = defaultdict(list)
    for x in satirlar:
        x["haftalik"] = yuvarla(x["satis"] * 7 / donem_gun)
        x["kapsama"] = yuvarla(x["stok"] / x["haftalik"]) if x["haftalik"] else None
        x["ihtiyac"] = x["fazla"] = 0
        if x["kapsama"] is not None and x["kapsama"] < alt:
            x["ihtiyac"] = math.ceil(hedef * x["haftalik"] - x["stok"])
        elif x["kapsama"] is None or x["kapsama"] > ust:
            x["fazla"] = max(0, math.floor(min(x["stok"] - hedef * x["haftalik"], x["stok"] - asgari)))
        x["durum"] = "Tükenme riski" if x["ihtiyac"] else "Fazla stok" if x["fazla"] else "Dengeli"
        x["kalan_ihtiyac"], x["kalan_fazla"] = x["ihtiyac"], x["fazla"]
        sku[tuple(katla(x[k]) for k in ("urun", "beden", "renk"))].append(x)
    transfer = []
    for grup in sku.values():
        for al in sorted((x for x in grup if x["ihtiyac"]), key=lambda x: (x["kapsama"], -x["ihtiyac"], x["magaza"])):
            for ver in sorted((x for x in grup if x["fazla"]), key=lambda v: (not (al["bolge"] and katla(v["bolge"]) == katla(al["bolge"])), -v["kalan_fazla"], v["magaza"])):
                adet = min(al["kalan_ihtiyac"], ver["kalan_fazla"])
                if adet < min_transfer or adet <= 0:
                    continue
                al["kalan_ihtiyac"] -= adet
                ver["kalan_fazla"] -= adet
                transfer.append({"urun": al["urun"], "beden": al["beden"], "renk": al["renk"], "gonderen": ver["magaza"], "alan": al["magaza"], "adet": adet,
                                 "ayni_bolge": bool(al["bolge"]) and katla(ver["bolge"]) == katla(al["bolge"]), "gonderen_stok": ver["stok"], "gonderen_kapsama": ver["kapsama"],
                                 "alan_stok": al["stok"], "alan_kapsama": al["kapsama"], "alan_haftalik": al["haftalik"]})
                if not al["kalan_ihtiyac"]:
                    break
    mag = {}
    for x in satirlar:
        m = mag.setdefault(katla(x["magaza"]), {"magaza": x["magaza"], "bolge": x["bolge"], "stok": Z, "risk": 0, "fazla_sku": 0, "gonderilen": 0, "alinan": 0, "kalan_ihtiyac": 0, "kalan_fazla": 0})
        m["stok"] += x["stok"]
        m["risk"] += bool(x["ihtiyac"])
        m["fazla_sku"] += bool(x["fazla"])
        m["gonderilen"] += x["fazla"] - x["kalan_fazla"]
        m["alinan"] += x["ihtiyac"] - x["kalan_ihtiyac"]
        m["kalan_ihtiyac"] += x["kalan_ihtiyac"]
        m["kalan_fazla"] += x["kalan_fazla"]
    return {"satirlar": satirlar, "transfer": sorted(transfer, key=lambda t: (t["gonderen"], t["alan"], t["urun"], t["beden"], t["renk"])),
            "eksik": sorted((x for x in satirlar if x["kalan_ihtiyac"]), key=lambda x: (x["kapsama"], x["magaza"])),
            "fazla": sorted((x for x in satirlar if x["kalan_fazla"]), key=lambda x: (-x["kalan_fazla"], x["magaza"])), "magaza": sorted(mag.values(), key=lambda m: m["magaza"]),
            "toplam": {"transfer": len(transfer), "adet": sum(t["adet"] for t in transfer), "ihtiyac": sum(x["ihtiyac"] for x in satirlar), "kalan_ihtiyac": sum(x["kalan_ihtiyac"] for x in satirlar),
                       "kalan_fazla": sum(x["kalan_fazla"] for x in satirlar)}, "esik": (donem_gun, hedef, alt, ust, asgari)}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    dg, hedef, alt, ust, asgari = r["esik"]
    ws = wb.active
    ws.title = "Transfer Önerileri"
    _baslik(ws, ["Gönderen Mağaza", "Alan Mağaza", "Ürün", "Beden", "Renk", "Transfer Adedi", "Aynı Bölge", "Gönderen Stok", "Gönderen Kapsama (hafta)", "Alan Stok", "Alan Kapsama (hafta)",
                 "Alan Haftalık Satış", "Onay (siz doldurun)"], (22, 22, 26, 8, 12, 10, 8, 10, 12, 9, 11, 11, 18))
    for t in r["transfer"]:
        _satir(ws, [t["gonderen"], t["alan"], t["urun"], t["beden"], t["renk"], t["adet"], "Evet" if t["ayni_bolge"] else "Hayır", t["gonderen_stok"], hafta(t["gonderen_kapsama"]) if t["gonderen_kapsama"] is None
                    else t["gonderen_kapsama"], t["alan_stok"], t["alan_kapsama"], t["alan_haftalik"], ""], dolgu={7: None if t["ayni_bolge"] else SARI})
    ws.append([])
    ws.append([f"Dönem {dg.normalize():f} gün. Alıcı: kapsama < {alt.normalize():f} hafta; verici: kapsama > {ust.normalize():f} hafta veya satışı yok; hedef {hedef.normalize():f} hafta; vericide bırakılan asgari "
               f"{asgari.normalize():f} adet. Eşleştirme sezgiseldir; nakliye maliyeti ve koli içi adet hesaba katılmaz."])
    ws = wb.create_sheet("Karşılanamayan İhtiyaç")
    _baslik(ws, ["Mağaza", "Ürün", "Beden", "Renk", "Stok", "Haftalık Satış", "Kapsama (hafta)", "İhtiyaç", "Transferle Karşılanan", "Kalan İhtiyaç"], (22, 26, 8, 12, 9, 10, 10, 9, 12, 10))
    for x in r["eksik"]:
        _satir(ws, [x["magaza"], x["urun"], x["beden"], x["renk"], x["stok"], x["haftalik"], x["kapsama"], x["ihtiyac"], x["ihtiyac"] - x["kalan_ihtiyac"], x["kalan_ihtiyac"]], dolgu={10: KIRMIZI})
    ws = wb.create_sheet("Kalan Fazla Stok")
    _baslik(ws, ["Mağaza", "Ürün", "Beden", "Renk", "Stok", "Haftalık Satış", "Kapsama (hafta)", "Fazla", "Transfer Edilen", "Kalan Fazla"], (22, 26, 8, 12, 9, 10, 12, 9, 10, 10))
    for x in r["fazla"]:
        _satir(ws, [x["magaza"], x["urun"], x["beden"], x["renk"], x["stok"], x["haftalik"], hafta(x["kapsama"]) if x["kapsama"] is None else x["kapsama"], x["fazla"], x["fazla"] - x["kalan_fazla"],
                    x["kalan_fazla"]], dolgu={10: SARI})
    ws = wb.create_sheet("Mağaza Özeti")
    _baslik(ws, ["Mağaza", "Bölge", "Toplam Stok", "Tükenme Riskli Kalem", "Fazla Stoklu Kalem", "Gönderilecek Adet", "Alınacak Adet", "Kalan İhtiyaç", "Kalan Fazla"], (22, 14, 10, 11, 10, 11, 10, 10, 10))
    for m in r["magaza"]:
        _satir(ws, [m["magaza"], m["bolge"], m["stok"], m["risk"], m["fazla_sku"], m["gonderilen"], m["alinan"], m["kalan_ihtiyac"], m["kalan_fazla"]], dolgu={8: KIRMIZI if m["kalan_ihtiyac"] else None})
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, **ayar) -> dict:
    satirlar, uy = oku(yol)
    r = oner(satirlar, **ayar)
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Fazla stoklu ve tükenme riskli mağazaları eşleştirip transfer önerir.")
    p.add_argument("--stok", type=Path, default=ORNEK / "magaza_stok_satis.csv", help="Mağaza, Bölge, Ürün, Beden, Renk, Stok, Satış (dönem adedi)")
    p.add_argument("--donem-gun", default="28", help="Satış verisinin kapsadığı gün sayısı (varsayılan 28)")
    p.add_argument("--hedef-hafta", default="4", help="Transfer sonrası hedeflenen kapsama (varsayılan 4 hafta)")
    p.add_argument("--min-hafta", default="2", help="Bu kapsamanın altı tükenme riski (varsayılan 2)")
    p.add_argument("--max-hafta", default="8", help="Bu kapsamanın üstü fazla stok (varsayılan 8)")
    p.add_argument("--asgari-stok", default="0", help="Verici mağazada bırakılacak asgari adet (varsayılan 0)")
    p.add_argument("--min-transfer", default="1", help="Bundan küçük transfer önerilmez (varsayılan 1)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "transfer_onerisi.xlsx")
    a = p.parse_args(argv)
    d = [para(x) for x in (a.donem_gun, a.hedef_hafta, a.min_hafta, a.max_hafta, a.asgari_stok, a.min_transfer)]
    if any(x is None or x < 0 for x in d) or d[0] == 0 or d[1] == 0:
        print("[X] --donem-gun ve --hedef-hafta pozitif; diğer parametreler sıfır veya pozitif sayı olmalı")
        return 2
    if not a.stok.exists():
        print(f"[X] Dosya bulunamadı: {a.stok}")
        return 1
    try:
        r = calistir(a.stok, a.cikti, donem_gun=d[0], hedef=d[1], alt=d[2], ust=d[3], asgari=d[4], min_transfer=d[5])
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {len(r['satirlar'])} satır · {t['transfer']} transfer önerisi, {t['adet']} adet · ihtiyaç {t['ihtiyac']} adet")
    if t["kalan_ihtiyac"]:
        print(f"[!] Transferle karşılanamayan ihtiyaç: {t['kalan_ihtiyac']} adet ({len(r['eksik'])} kalem)")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
