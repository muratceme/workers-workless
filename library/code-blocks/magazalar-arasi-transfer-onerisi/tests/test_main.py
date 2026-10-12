import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def s(magaza, stok, satis, bolge="", urun="T", beden="M", renk=""):
    return {"magaza": magaza, "bolge": bolge, "urun": urun, "beden": beden, "renk": renk, "stok": D(stok), "satis": D(satis)}


def tek(stok, satis, **kw):
    return main.oner([s("A", stok, satis)], **kw)["satirlar"][0]


class OneriTesti(unittest.TestCase):
    def test_alici_ve_verici(self):
        x = tek(2, 16)
        self.assertEqual((x["haftalik"], x["kapsama"], x["ihtiyac"], x["fazla"], x["durum"]), (D("4.00"), D("0.50"), 14, 0, "Tükenme riski"))
        x = tek(40, 8)
        self.assertEqual((x["kapsama"], x["ihtiyac"], x["fazla"], x["durum"]), (D("20.00"), 0, 32, "Fazla stok"))
        x = tek(30, 0)
        self.assertEqual((x["kapsama"], x["fazla"]), (None, 30))
        self.assertEqual([tek(st, 4)["durum"] for st in (1, 2, 8, 9)], ["Tükenme riski", "Dengeli", "Dengeli", "Fazla stok"])        # tam 2 ve tam 8 hafta dengeli
        self.assertEqual((tek(30, 0, asgari=D(4))["fazla"], tek(9, 4, asgari=D(6))["fazla"], tek(2, 0, asgari=D(5))["fazla"]), (26, 3, 0))
        self.assertEqual(tek(3, 10, donem_gun=D(30))["ihtiyac"], 7)                                  # haftalık 2,33 → 4 × 2,33 − 3 = 6,32 → 7
        with self.assertRaises(ValueError):
            main.oner([])
        with self.assertRaises(ValueError):
            main.oner([s("A", 1, 1)], hedef=D(9))

    def test_eslestirme(self):
        l = [s("Alıcı", 2, 16, "X"), s("Uzak", 50, 0, "Y"), s("Yakın", 10, 0, "X")]
        self.assertEqual([(t["gonderen"], t["adet"], t["ayni_bolge"]) for t in main.oner(l)["transfer"]], [("Uzak", 4, False), ("Yakın", 10, True)])   # önce aynı bölge, kalan uzaktan
        l = [s("Acil", 1, 12), s("Az acil", 3, 20), s("Verici", 20, 4)]
        r = main.oner(l)
        self.assertEqual({t["alan"]: t["adet"] for t in r["transfer"]}, {"Acil": 11, "Az acil": 5})
        self.assertEqual([(x["magaza"], x["kalan_ihtiyac"]) for x in r["eksik"]], [("Az acil", 12)])
        self.assertEqual(main.oner([s("A", 1, 4), s("V", 6, 0)], min_transfer=D(5))["transfer"], [])                 # ihtiyaç 3 < en küçük transfer
        self.assertEqual(main.oner([s("A", 1, 4), s("V", 6, 0, beden="L")])["transfer"], [])                         # başka beden eşleşmez
        r = main.oner([s("A", 2, 16), s("V1", 5, 0), s("V2", 5, 0), s("V3", 9, 0)])
        self.assertEqual([(t["gonderen"], t["adet"]) for t in r["transfer"]], [("V1", 5), ("V3", 9)])                # fazlası en çok olandan başlar
        self.assertEqual((r["toplam"]["adet"], r["toplam"]["kalan_fazla"]), (14, 5))


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "magaza_stok_satis.csv", cikti)
            self.assertEqual([(t["gonderen"], t["alan"], t["urun"], t["beden"], t["adet"], t["ayni_bolge"]) for t in r["transfer"]],
                             [("Beşiktaş", "Bakırköy", "Slim Jean", "32", 8, True), ("Ümraniye", "Beşiktaş", "Basic Tişört", "L", 11, False), ("Ümraniye", "Kadıköy", "Basic Tişört", "L", 5, True),
                              ("Ümraniye", "Kadıköy", "Basic Tişört", "M", 14, True)])
            self.assertEqual(r["toplam"], {"transfer": 4, "adet": 38, "ihtiyac": 50, "kalan_ihtiyac": 12, "kalan_fazla": 91})
            self.assertEqual({m["magaza"]: (m["gonderilen"], m["alinan"], m["kalan_ihtiyac"], m["kalan_fazla"]) for m in r["magaza"]},
                             {"Kadıköy": (0, 19, 12, 30), "Ümraniye": (30, 0, 0, 18), "Beşiktaş": (8, 11, 0, 13), "Bakırköy": (0, 8, 0, 30)})
            self.assertEqual(load_workbook(cikti).sheetnames, ["Transfer Önerileri", "Karşılanamayan İhtiyaç", "Kalan Fazla Stok", "Mağaza Özeti", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "s.csv"
            y.write_text("Mağaza;Ürün;Stok;Satış\nA;T;2;16\nA;T;5;1\nB;T;yok;3\nC;T;1.000;0\n;T;1;1\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual([(x["magaza"], x["stok"], x["beden"]) for x in l], [("A", D(2), ""), ("C", D(1000), "")])
            self.assertEqual(sorted(x["tur"] for x in uy), ["Değer okunamadı", "Mükerrer satır"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--stok", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--donem-gun", "14", "--hedef-hafta", "3", "--asgari-stok", "2"]), 0)
            self.assertEqual(main.main(arg + ["--donem-gun", "0"]), 2)
            self.assertEqual(main.main(arg + ["--min-hafta", "5"]), 1)
            self.assertEqual(main.main(["--stok", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Mağaza;Ürün;Stok;Satış\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
