import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

ORNEK = main.ORNEK


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "f.xlsx"
        cls.s = main.calistir(ORNEK / "satirlar.csv", cls.cikti, ORNEK / "faturalar.csv")
        cls.f = {f.no[-3:]: f for f in cls.s["faturalar"]}
        cls.b = {(b["fatura"][-3:], b["tur"]) for b in cls.s["bulgular"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_satir_hatalari(self):
        self.assertTrue({("102", "Tutar hesabı"), ("105", "Geçersiz KDV oranı"), ("106", "Geçersiz tevkifat oranı"), ("107", "Tevkifat hesabı"),
                         ("108", "Oran 0, KDV var")} <= self.b)
        self.assertEqual(len([b for b in self.s["bulgular"] if b["onem"] == "Hata"]), 5)

    def test_fatura_durumu(self):
        d = {k: f.durum for k, f in self.f.items()}
        self.assertEqual([k for k, v in sorted(d.items()) if v == "Tamam"], ["101", "103", "104"])
        self.assertEqual(d["109"], "Yuvarlama")                       # toplam KDV 0,03 TL farklı
        self.assertIn(("109", "Yuvarlama farkı"), self.b)
        self.assertEqual(self.f["101"].satir_tutar, D("6880.00"))     # dosyanın sonundaki 3. satır da faturaya eklendi

    def test_excel(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Bulgular", "Fatura Özeti", "Satırlar", "KDV Oran Özeti"])
        self.assertEqual(wb["Satırlar"].max_row, 18)
        self.assertEqual(wb["KDV Oran Özeti"]["A6"].value, "%20")      # %0, %1, %10, %18, %20


class KuralTesti(unittest.TestCase):
    def test_yardimcilar(self):
        self.assertEqual(main.tevkifat_orani("9/10"), D("0.9"))
        self.assertEqual(main.tevkifat_orani("%50"), D("0.5"))
        self.assertEqual(main.tevkifat_orani("0,7"), D("0.7"))
        self.assertEqual(main.karsilastir(D("10.00"), D("10.03"), D("0.01"), D("0.05")), "Bilgi")
        self.assertIsNone(main.karsilastir(D("10.00"), D("10.01"), D("0.01"), D("0.05")))

    def test_eski_oran_ve_efatura_ciktisi(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp) / "e.xlsx"
            wb = Workbook()
            fa = wb.active
            fa.title = "Faturalar"
            fa.append(["Fatura No", "Satıcı", "Mal/Hizmet Toplamı", "Toplam KDV", "KDV Tevkifatı", "Vergiler Dahil", "Ödenecek"])
            fa.append(["X1", "A", 1000, 180, 0, 1180, 1180])
            fa.append(["X2", "A", 500, 100, 0, 600, 650])
            sa = wb.create_sheet("Satırlar")
            sa.append(["Fatura No", "Tarih", "Satıcı", "Alıcı", "Sıra", "Mal/Hizmet", "Miktar", "Birim", "Birim Fiyat", "İskonto", "Tutar", "KDV Oranı", "KDV Tutarı"])
            sa.append(["X1", "15.03.2023", "A", "B", 1, "Hizmet", 1, "Adet", 1000, 0, 1000, 18.0, 180])
            sa.append(["X2", "15.03.2024", "A", "B", 1, "Mal", 2, "Adet", 250, 0, 500, 20.0, 100])
            sa.append(["X2", "15.03.2024", "A", "B", 1, "Mal", 2, "Adet", 250, 0, 500, 20.0, 100])
            wb.save(t)
            s = main.calistir(t, Path(tmp) / "o.xlsx")
            f = {x.no: x for x in s["faturalar"]}
            self.assertEqual(f["X1"].durum, "Tamam")                  # 2023 Mart'ta %18 geçerli
            turler = {b["tur"] for b in s["bulgular"] if b["fatura"] == "X2"}
            self.assertTrue({"Mükerrer satır", "Mal/hizmet toplamı", "Ödenecek tutar"} <= turler)

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--satirlar", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
