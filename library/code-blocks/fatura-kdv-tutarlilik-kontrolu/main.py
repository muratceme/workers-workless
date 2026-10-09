"""
Fatura KDV Tutarlılık Kontrolü — Workers / Workless kod bloğu
Muhasebe › Muhasebe Elemanı

Fatura satırlarında ve fatura toplamlarında matrah, KDV oranı, KDV tutarı ve tevkifat tutarlılığını kontrol eder:
  - Satır: tutar = miktar × birim fiyat − iskonto; KDV = tutar × oran; oran fatura tarihinde geçerli mi
    (10.07.2023 ve sonrası %0, 1, 10, 20; öncesi %0, 1, 8, 18); oran 0 iken KDV, oran > 0 iken KDV 0.
  - Tevkifat: oran 2/10, 3/10, 4/10, 5/10, 7/10, 9/10 veya 10/10 mu; tevkif edilen KDV = KDV × tevkifat oranı.
  - Fatura toplamı (isteğe bağlı fatura dosyası): mal/hizmet toplamı, oran bazında KDV (matrah × oran),
    tevkifat, vergiler dahil tutar ve ödenecek = vergiler dahil − tevkifat.
  - Fark ≤ tolerans → tamam; ≤ yuvarlama sınırı → "Yuvarlama farkı" (bilgi); üstü → hata.
  - Aynı fatura no + satıcı + sıra tekrarı.
Girdi: tek tablo (satırlar) + isteğe bağlı fatura başlık tablosu; "E-Fatura Okuma ve Listeleme" paketinin
Excel çıktısı (Faturalar + Satırlar sayfaları) doğrudan verilebilir. İnternete bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: 9 fatura, 17 satır
    python main.py --satirlar satirlar.xlsx --faturalar faturalar.xlsx
    python main.py --satirlar e_fatura_listesi.xlsx                # e-fatura paketinin çıktısı
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
K2 = Decimal("0.01")
ORAN_DEGISIMI = date(2023, 7, 10)            # Cumhurbaşkanı Kararı 7346 (RG 07.07.2023): %8 → %10, %18 → %20
ORANLAR_YENI = {Decimal(0), Decimal(1), Decimal(10), Decimal(20)}
ORANLAR_ESKI = {Decimal(0), Decimal(1), Decimal(8), Decimal(18)}
TEVKIFAT_ORANLARI = {Decimal(n) / 10 for n in (2, 3, 4, 5, 7, 9, 10)}

SATIR_SUTUNLARI = {"no": ("fatura no", "belge no", "fatura numarasi"), "tarih": ("tarih", "fatura tarihi"), "satici": ("satici", "satici unvan", "tedarikci"),
                   "sira": ("sira", "satir no", "sira no"), "urun": ("mal hizmet", "urun", "aciklama"), "miktar": ("miktar",),
                   "fiyat": ("birim fiyat",), "iskonto": ("iskonto", "indirim"), "tutar": ("tutar", "matrah", "satir tutari", "mal hizmet tutari"),
                   "oran": ("kdv orani", "kdv"), "kdv": ("kdv tutari",), "t_oran": ("tevkifat orani",), "t_tutar": ("tevkif edilen kdv", "tevkifat tutari")}
FATURA_SUTUNLARI = {"no": ("fatura no", "belge no"), "satici": ("satici", "satici unvan"), "tarih": ("tarih",),
                    "mal": ("mal hizmet toplami", "mal hizmet toplam tutari", "matrah toplami"), "kdv": ("toplam kdv", "hesaplanan kdv"),
                    "tevkifat": ("kdv tevkifati", "tevkifat", "tevkif edilen kdv"), "dahil": ("vergiler dahil", "vergiler dahil toplam"),
                    "odenecek": ("odenecek", "odenecek tutar")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("%", "").replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tevkifat_orani(x) -> Decimal | None:
    """'9/10', '90', '%90', 0,9 → 0.9"""
    if x in (None, ""):
        return None
    s = str(x).strip().replace("%", "")
    m = re.fullmatch(r"(\d+)\s*/\s*(\d+)", s)
    if m:
        return Decimal(m[1]) / Decimal(m[2]) if int(m[2]) else None
    v = para(s)
    if v is None:
        return None
    return v / 100 if v > 1 else v


def tl(x: Decimal | None) -> str:
    return "—" if x is None else f"{x.quantize(K2, ROUND_HALF_UP):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def kesir(o: Decimal) -> str:
    return f"{(o * 10).normalize():f}/10"


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def yuvarla(x: Decimal) -> Decimal:
    return x.quantize(K2, ROUND_HALF_UP)


def tablo_oku(yol: Path, sayfa: str | None = None) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            ws = wb[sayfa] if sayfa and sayfa in wb.sheetnames else wb.active
            satirlar = [list(r) for r in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(s: list[list], ad: str, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{ad}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


def sayfa_var(yol: Path, ad: str) -> bool:
    if yol.suffix.lower() not in {".xlsx", ".xlsm"}:
        return False
    wb = load_workbook(yol, read_only=True)
    try:
        return ad in wb.sheetnames
    finally:
        wb.close()


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Satir:
    no: str
    tarih: date | None
    satici: str
    sira: str
    urun: str
    miktar: Decimal | None
    fiyat: Decimal | None
    iskonto: Decimal
    tutar: Decimal | None
    oran: Decimal | None
    kdv: Decimal | None
    t_oran: Decimal | None
    t_tutar: Decimal | None
    bulgular: list = field(default_factory=list)      # (önem, tür, açıklama)


@dataclass
class Fatura:
    no: str
    satici: str
    tarih: date | None
    satirlar: list
    mal: Decimal | None = None
    kdv: Decimal | None = None
    tevkifat: Decimal | None = None
    dahil: Decimal | None = None
    odenecek: Decimal | None = None
    baslik_var: bool = False
    bulgular: list = field(default_factory=list)

    @property
    def satir_tutar(self) -> Decimal:
        return sum((s.tutar or Decimal(0) for s in self.satirlar), Decimal(0))

    @property
    def satir_kdv(self) -> Decimal:
        return sum((s.kdv or Decimal(0) for s in self.satirlar), Decimal(0))

    @property
    def satir_tevkifat(self) -> Decimal:
        return sum((s.t_tutar or Decimal(0) for s in self.satirlar), Decimal(0))

    def oran_bazinda(self) -> dict[Decimal, list[Decimal]]:
        g = defaultdict(lambda: [Decimal(0), Decimal(0)])
        for s in self.satirlar:
            if s.oran is not None and s.tutar is not None:
                g[s.oran][0] += s.tutar
                g[s.oran][1] += s.kdv or Decimal(0)
        return dict(g)

    @property
    def durum(self) -> str:
        onemler = {b[0] for b in self.bulgular} | {b[0] for s in self.satirlar for b in s.bulgular}
        return "Hata" if "Hata" in onemler else "Yuvarlama" if "Bilgi" in onemler else "Tamam"


def oku(satir_yolu: Path, fatura_yolu: Path | None) -> tuple[list[Fatura], list[dict]]:
    uyarilar = []
    efatura = sayfa_var(satir_yolu, "Satırlar")
    satirlar = []
    for r in kayitlar(tablo_oku(satir_yolu, "Satırlar" if efatura else None), satir_yolu.name, SATIR_SUTUNLARI, ("no", "tutar")):
        if not metin(r.get("no")):
            continue
        satirlar.append(Satir(metin(r["no"]), tarih(r.get("tarih")), metin(r.get("satici")), metin(r.get("sira")) or str(r["_satir"]), metin(r.get("urun")),
                              para(r.get("miktar")), para(r.get("fiyat")), para(r.get("iskonto")) or Decimal(0), para(r.get("tutar")),
                              None if r.get("oran") in (None, "") else (lambda v: v * 100 if v is not None and 0 < v < 1 else v)(para(r.get("oran"))),
                              para(r.get("kdv")), tevkifat_orani(r.get("t_oran")), para(r.get("t_tutar"))))
    faturalar: dict[tuple[str, str], Fatura] = {}
    for s in satirlar:
        f = faturalar.setdefault((s.no, katla(s.satici)), Fatura(s.no, s.satici, s.tarih, []))
        f.satirlar.append(s)
    if efatura and not fatura_yolu:
        fatura_yolu, sayfa = satir_yolu, "Faturalar"
    else:
        sayfa = None
    if fatura_yolu:
        for r in kayitlar(tablo_oku(fatura_yolu, sayfa), fatura_yolu.name, FATURA_SUTUNLARI, ("no",)):
            no = metin(r.get("no"))
            if not no:
                continue
            adaylar = [f for (n, sa), f in faturalar.items() if n == no and (not r.get("satici") or sa == katla(r.get("satici")) or not sa)]
            if not adaylar:
                uyarilar.append({"onem": "Uyarı", "tur": "Satırı olmayan fatura", "fatura": no, "aciklama": "Fatura başlık dosyasında var, satır dosyasında yok"})
                continue
            f = adaylar[0]
            f.mal, f.kdv, f.tevkifat, f.dahil, f.odenecek = (para(r.get(k)) for k in ("mal", "kdv", "tevkifat", "dahil", "odenecek"))
            f.tarih = f.tarih or tarih(r.get("tarih"))
            f.baslik_var = True
    return list(faturalar.values()), uyarilar


# ----------------------------------------------------------------------------
# Kontroller
# ----------------------------------------------------------------------------

def karsilastir(beklenen: Decimal, beyan: Decimal, tol: Decimal, yuv: Decimal) -> str | None:
    """None: tamam · 'Bilgi': yuvarlama farkı · 'Hata'"""
    f = abs(beklenen - beyan)
    return None if f <= tol else "Bilgi" if f <= yuv else "Hata"


def kontrol_et(faturalar: list[Fatura], tol: Decimal = K2, yuv: Decimal = Decimal("0.05")) -> list[dict]:
    uyarilar = []
    for f in faturalar:
        sayac = Counter(s.sira for s in f.satirlar)
        for sira, n in sayac.items():
            if n > 1:
                f.bulgular.append(("Hata", "Mükerrer satır", f"Sıra {sira} {n} kez var"))
        for s in f.satirlar:
            b = s.bulgular
            if s.tutar is None:
                b.append(("Hata", "Tutar okunamadı", "Satır tutarı boş veya sayı değil"))
                continue
            if s.tutar < 0 or (s.miktar is not None and s.miktar <= 0):
                b.append(("Uyarı", "Negatif / sıfır değer", f"Miktar {s.miktar}, tutar {tl(s.tutar)}; iade faturası olabilir"))
            if s.miktar is not None and s.fiyat is not None:
                hesap = yuvarla(s.miktar * s.fiyat - s.iskonto)
                d = karsilastir(hesap, s.tutar, tol, yuv)
                if d:
                    b.append((d, "Tutar hesabı" if d == "Hata" else "Yuvarlama farkı",
                              f"{s.miktar:g} × {tl(s.fiyat)}" + (f" − {tl(s.iskonto)}" if s.iskonto else "") + f" = {tl(hesap)}; satırda {tl(s.tutar)}"
                              f" (fark {tl(s.tutar - hesap)})"))
            if s.oran is None:
                b.append(("Hata", "KDV oranı yok", "KDV oranı boş"))
                continue
            gecerli = ORANLAR_YENI if not s.tarih or s.tarih >= ORAN_DEGISIMI else ORANLAR_ESKI
            if s.oran not in gecerli:
                b.append(("Hata", "Geçersiz KDV oranı", f"%{s.oran:g} fatura tarihinde ({s.tarih:%d.%m.%Y}) geçerli değil; geçerli oranlar: "
                          + ", ".join(f"%{o:g}" for o in sorted(gecerli)) if s.tarih else f"%{s.oran:g} geçerli bir oran değil"))
            if s.kdv is None:
                b.append(("Hata", "KDV tutarı yok", "KDV tutarı boş"))
            elif s.oran == 0 and s.kdv != 0:
                b.append(("Hata", "Oran 0, KDV var", f"KDV oranı %0 ama KDV tutarı {tl(s.kdv)}"))
            elif s.oran > 0 and s.kdv == 0:
                b.append(("Hata", "KDV eksik", f"KDV oranı %{s.oran:g} ama KDV tutarı 0"))
            else:
                beklenen = yuvarla(s.tutar * s.oran / 100)
                d = karsilastir(beklenen, s.kdv, tol, yuv)
                if d:
                    b.append((d, "KDV hesabı" if d == "Hata" else "Yuvarlama farkı", f"{tl(s.tutar)} × %{s.oran:g} = {tl(beklenen)}; satırda {tl(s.kdv)}"))
            if s.t_oran is not None or s.t_tutar:
                if s.t_oran is None:
                    b.append(("Hata", "Tevkifat oranı yok", f"Tevkif edilen KDV {tl(s.t_tutar)} var ama oran yok"))
                elif s.t_oran not in TEVKIFAT_ORANLARI:
                    b.append(("Hata", "Geçersiz tevkifat oranı", f"{kesir(s.t_oran)} tanımlı oranlardan değil (2, 3, 4, 5, 7, 9 veya 10 / 10)"))
                elif s.kdv is not None:
                    beklenen = yuvarla(s.kdv * s.t_oran)
                    d = karsilastir(beklenen, s.t_tutar or Decimal(0), tol, yuv)
                    if d:
                        b.append((d, "Tevkifat hesabı" if d == "Hata" else "Yuvarlama farkı",
                                  f"{tl(s.kdv)} × {kesir(s.t_oran)} = {tl(beklenen)}; satırda {tl(s.t_tutar)}"))
        if not f.baslik_var:
            continue
        n = max(len(f.satirlar), 1)
        ytol = max(yuv, K2 * n)          # toplamlarda satır başına bir kuruş yuvarlama payı
        if f.mal is not None:
            d = karsilastir(f.satir_tutar, f.mal, tol, ytol)
            if d:
                f.bulgular.append((d, "Mal/hizmet toplamı", f"Satırlar toplamı {tl(f.satir_tutar)}; faturada {tl(f.mal)}"))
        if f.kdv is not None:
            oran_kdv = sum((yuvarla(m * o / 100) for o, (m, _) in f.oran_bazinda().items()), Decimal(0))
            d1, d2 = karsilastir(f.satir_kdv, f.kdv, tol, ytol), karsilastir(oran_kdv, f.kdv, tol, ytol)
            if d1 and d2:
                d = "Hata" if "Hata" in (d1, d2) and min(abs(f.satir_kdv - f.kdv), abs(oran_kdv - f.kdv)) > ytol else "Bilgi"
                f.bulgular.append((d, "Toplam KDV" if d == "Hata" else "Yuvarlama farkı",
                                   f"Satır KDV toplamı {tl(f.satir_kdv)}, oran bazında matrah × oran {tl(oran_kdv)}; faturada {tl(f.kdv)}"))
        if f.tevkifat is not None and (f.tevkifat or f.satir_tevkifat):
            d = karsilastir(f.satir_tevkifat, f.tevkifat, tol, ytol)
            if d:
                f.bulgular.append((d, "Tevkifat toplamı", f"Satırlarda tevkif edilen {tl(f.satir_tevkifat)}; faturada {tl(f.tevkifat)}"))
        if f.dahil is not None and f.mal is not None and f.kdv is not None:
            d = karsilastir(f.mal + f.kdv, f.dahil, tol, ytol)
            if d:
                f.bulgular.append((d, "Vergiler dahil tutar", f"Mal/hizmet {tl(f.mal)} + KDV {tl(f.kdv)} = {tl(f.mal + f.kdv)}; faturada {tl(f.dahil)}"))
        if f.odenecek is not None and f.dahil is not None:
            d = karsilastir(f.dahil - (f.tevkifat or 0), f.odenecek, tol, ytol)
            if d:
                f.bulgular.append((d, "Ödenecek tutar", f"Vergiler dahil {tl(f.dahil)} − tevkifat {tl(f.tevkifat or Decimal(0))} = "
                                   f"{tl(f.dahil - (f.tevkifat or 0))}; faturada {tl(f.odenecek)}"))
    return uyarilar


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Hata": "FDE2E1", "Uyarı": "FFF4CE", "Bilgi": "E8F0FE", "Tamam": "E3F4E1", "Yuvarlama": "E8F0FE"}
KONTROL = PatternFill("solid", fgColor="FFF4CE")
UST = Alignment(vertical="top", wrap_text=True)
PF = "#,##0.00"


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def _f(x):
    return None if x is None else float(x)


def bulgu_listesi(faturalar: list[Fatura]) -> list[dict]:
    sira = {"Hata": 0, "Uyarı": 1, "Bilgi": 2}
    lst = [{"onem": o, "tur": t, "fatura": f.no, "satici": f.satici, "sira": "", "aciklama": a} for f in faturalar for o, t, a in f.bulgular]
    lst += [{"onem": o, "tur": t, "fatura": f.no, "satici": f.satici, "sira": s.sira, "aciklama": a} for f in faturalar for s in f.satirlar
            for o, t, a in s.bulgular]
    return sorted(lst, key=lambda b: (sira[b["onem"]], b["fatura"], str(b["sira"]).zfill(4)))


def rapor_yaz(cikti: Path, faturalar: list[Fatura], bulgular: list[dict]) -> None:
    wb = Workbook()
    bl = wb.active
    bl.title = "Bulgular"
    _baslik(bl, ["Önem", "Tür", "Fatura No", "Satıcı", "Satır", "Açıklama", "Karar / Düzeltme"], (8, 22, 19, 26, 6, 80, 24))
    for b in bulgular:
        bl.append([b["onem"], b["tur"], b["fatura"], b["satici"], b["sira"], b["aciklama"], ""])
        bl.cell(bl.max_row, 1).fill = PatternFill("solid", fgColor=RENK[b["onem"]])
        bl.cell(bl.max_row, 6).alignment = UST
        bl.cell(bl.max_row, 7).fill = KONTROL

    fo = wb.create_sheet("Fatura Özeti")
    _baslik(fo, ["Fatura No", "Tarih", "Satıcı", "Satır", "Satırlar Toplamı", "Faturada Mal/Hizmet", "Satır KDV", "Faturada KDV", "Satır Tevkifat",
                 "Faturada Tevkifat", "Vergiler Dahil", "Ödenecek", "Durum", "Bulgu"], (19, 11, 26, 6, 14, 14, 12, 12, 12, 12, 14, 14, 10, 7))
    for f in faturalar:
        n = len(f.bulgular) + sum(len(s.bulgular) for s in f.satirlar)
        fo.append([f.no, f.tarih, f.satici, len(f.satirlar), float(f.satir_tutar), _f(f.mal), float(f.satir_kdv), _f(f.kdv), float(f.satir_tevkifat),
                   _f(f.tevkifat), _f(f.dahil), _f(f.odenecek), f.durum, n])
        r = fo.max_row
        fo.cell(r, 2).number_format = "DD.MM.YYYY"
        for j in range(5, 13):
            fo.cell(r, j).number_format = PF
        fo.cell(r, 13).fill = PatternFill("solid", fgColor=RENK[f.durum])
    fo.auto_filter.ref = f"A1:N{fo.max_row}"

    st = wb.create_sheet("Satırlar")
    _baslik(st, ["Fatura No", "Sıra", "Mal/Hizmet", "Miktar", "Birim Fiyat", "İskonto", "Tutar", "Hesaplanan Tutar", "KDV Oranı", "KDV", "Hesaplanan KDV",
                 "Tevkifat", "Tevkif Edilen", "Hesaplanan Tevkifat", "Bulgu"], (19, 6, 26, 8, 11, 9, 12, 12, 8, 11, 11, 8, 11, 11, 40))
    for f in faturalar:
        for s in f.satirlar:
            h_tutar = yuvarla(s.miktar * s.fiyat - s.iskonto) if s.miktar is not None and s.fiyat is not None else None
            h_kdv = yuvarla(s.tutar * s.oran / 100) if s.tutar is not None and s.oran is not None else None
            h_t = yuvarla(s.kdv * s.t_oran) if s.kdv is not None and s.t_oran is not None else None
            st.append([f.no, s.sira, s.urun, _f(s.miktar), _f(s.fiyat), _f(s.iskonto), _f(s.tutar), _f(h_tutar), _f(s.oran), _f(s.kdv), _f(h_kdv),
                       kesir(s.t_oran) if s.t_oran is not None else "", _f(s.t_tutar), _f(h_t), "; ".join(t for _, t, _ in s.bulgular)])
            r = st.max_row
            for j in (5, 6, 7, 8, 10, 11, 13, 14):
                st.cell(r, j).number_format = PF
            if s.bulgular:
                st.cell(r, 15).fill = PatternFill("solid", fgColor=RENK[min((b[0] for b in s.bulgular), key=["Hata", "Uyarı", "Bilgi"].index)])

    ko = wb.create_sheet("KDV Oran Özeti")
    _baslik(ko, ["KDV Oranı", "Matrah", "Satırlardaki KDV", "Matrah × Oran", "Fark"], (10, 16, 16, 16, 12))
    g = defaultdict(lambda: [Decimal(0), Decimal(0)])
    for f in faturalar:
        for o, (m, k) in f.oran_bazinda().items():
            g[o][0] += m
            g[o][1] += k
    for o, (m, k) in sorted(g.items()):
        h = yuvarla(m * o / 100)
        ko.append([f"%{o:g}", float(m), float(k), float(h), float(k - h)])
        for j in (2, 3, 4, 5):
            ko.cell(ko.max_row, j).number_format = PF
    ko.append([])
    ko.append(["Oranlar: 10.07.2023 ve sonrası %1, %10, %20 (Cumhurbaşkanı Kararı 7346, RG 07.07.2023); öncesi %1, %8, %18. "
               "Ürüne hangi oranın uygulanacağı bu pakette kontrol edilmez."])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(satir_yolu: Path, cikti: Path, fatura_yolu: Path | None = None, tol: Decimal = K2, yuv: Decimal = Decimal("0.05")) -> dict:
    faturalar, uyarilar = oku(satir_yolu, fatura_yolu)
    if not faturalar:
        raise ValueError(f"{satir_yolu.name}: fatura satırı bulunamadı")
    uyarilar += kontrol_et(faturalar, tol, yuv)
    bulgular = bulgu_listesi(faturalar) + uyarilar
    rapor_yaz(cikti, faturalar, bulgular)
    return {"faturalar": faturalar, "bulgular": bulgular}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Fatura satırlarında matrah, KDV oranı, KDV tutarı ve tevkifat tutarlılığını kontrol eder.")
    p.add_argument("--satirlar", type=Path, default=ORNEK / "satirlar.csv",
                   help="Fatura No, Tarih, Satıcı, Sıra, Mal/Hizmet, Miktar, Birim Fiyat, İskonto, Tutar, KDV Oranı, KDV Tutarı, [Tevkifat Oranı, Tevkif Edilen KDV]")
    p.add_argument("--faturalar", type=Path, help="İsteğe bağlı: Fatura No, Satıcı, Mal/Hizmet Toplamı, Toplam KDV, KDV Tevkifatı, Vergiler Dahil, Ödenecek")
    p.add_argument("--tolerans", type=float, default=0.01, help="Bu farka kadar tamam sayılır, TL (varsayılan 0,01)")
    p.add_argument("--yuvarlama", type=float, default=0.05, help="Bu farka kadar 'yuvarlama farkı' sayılır, TL (varsayılan 0,05)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "fatura_kdv_kontrolu.xlsx")
    a = p.parse_args(argv)
    fat = a.faturalar or (ORNEK / "faturalar.csv" if a.satirlar == ORNEK / "satirlar.csv" else None)
    for y in (a.satirlar, fat):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        s = calistir(a.satirlar, a.cikti, fat, Decimal(str(a.tolerans)), Decimal(str(a.yuvarlama)))
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    d = Counter(f.durum for f in s["faturalar"])
    print(f"[OK] {len(s['faturalar'])} fatura, {sum(len(f.satirlar) for f in s['faturalar'])} satır · " + " · ".join(f"{k} {v}" for k, v in d.items()))
    for b in s["bulgular"]:
        if b["onem"] == "Hata":
            print(f"[!] {b['fatura']}" + (f" satır {b['sira']}" if b["sira"] else "") + f" {b['tur']}: {b['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
