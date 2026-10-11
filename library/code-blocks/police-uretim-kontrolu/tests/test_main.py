import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def belge(no, teklif="", teminat=None, sart="", sigortali="Firma", bas=date(2026, 11, 1), bit=date(2027, 11, 1)):
    tem = teminat if teminat is not None else {"Yangın": (1000, 10, "")}
    return {"no": no, "teklif": teklif, "sigortali": sigortali, "bas": bas, "bit": bit, "sart": main.sartlar(sart),
            "teminat": {main.katla(k): {"ad": k, "bedel": None if v[0] is None else D(str(v[0])), "prim": None if v[1] is None else D(str(v[1])), "muafiyet": v[2]} for k, v in tem.items()}}


def calis(t, p, **kw):
    r = main.karsilastir({"T1": t} if t else {}, {x["no"]: x for x in (p if isinstance(p, list) else [p])}, **kw)
    return r, [(f["duzey"], f["konu"]) for f in r["farklar"]]


class KarsilastirmaTesti(unittest.TestCase):
    def test_uyumlu_ve_tolerans(self):
        r, f = calis(belge("T1", sart="Kloz A | Kloz B"), belge("P1", "T1", {"YANGIN": (1000.5, 11, "")}, sart="kloz b;KLOZ A"))
        self.assertEqual((f, r["ozet"][0]["sonuc"], r["ozet"][0]["teklif_prim"], r["ozet"][0]["police_prim"]), ([], "Uyumlu", D(10), D(11)))
        _, f = calis(belge("T1"), belge("P1", "T1", {"Yangın": (1001.01, 11.01, "")}))
        self.assertEqual(f, [("Yüksek", "Bedel farklı"), ("Yüksek", "Prim farklı")])
        self.assertEqual(calis(belge("T1"), belge("P1", "T1", {"Yangın": (1001.01, 11.01, "")}), tolerans=D(2))[1], [])
        _, f = calis(belge("T1"), belge("P1", "T1", {"Yangın": (1000, None, "")}))
        self.assertEqual(f, [("Yüksek", "Prim farklı")])

    def test_teminat_muafiyet_sart(self):
        t = belge("T1", teminat={"Yangın": (1000, 10, ""), "Deprem": (1000, 20, "%2"), "Hırsızlık": (100, 5, "")}, sart="Sprinkler klozu")
        p = belge("P1", "T1", {"Yangın": (1000, 10, ""), "Deprem": (1000, 20, "%5"), "Cam": (50, 1, "")}, sart="Enflasyon klozu")
        r, f = calis(t, p)
        self.assertEqual(f, [("Orta", "Muafiyet farklı"), ("Yüksek", "Teminat poliçede yok"), ("Yüksek", "Teminat teklifte yok"), ("Yüksek", "Özel şart poliçede yok"),
                             ("Orta", "Özel şart teklifte yok")])
        self.assertEqual((r["farklar"][1]["teminat"], r["farklar"][1]["teklifte"], r["farklar"][0]["policede"]), ("Hırsızlık", "bedel 100,00, prim 5,00", "%5"))
        self.assertEqual((r["ozet"][0]["sonuc"], r["ozet"][0]["fark"], r["ozet"][0]["teklif_prim"], r["ozet"][0]["police_prim"]), ("Fark var", 5, D(35), D(31)))

    def test_baslik_alanlari(self):
        _, f = calis(belge("T1"), belge("P1", "T1", sigortali="FİRMA A.Ş.", bas=date(2026, 11, 2), bit=date(2027, 11, 2)))
        self.assertEqual(f, [("Orta", "Sigortalı adı farklı"), ("Yüksek", "Başlangıç tarihi farklı"), ("Yüksek", "Bitiş tarihi farklı")])
        self.assertEqual(calis(belge("T1"), belge("P1", "T1", sigortali="firma", bas=None))[1], [])        # boş tarih karşılaştırılmaz

    def test_teklif_yok_mukerrer_uretilmemis(self):
        r, f = calis(belge("T1"), [belge("P1", "T1"), belge("P2", "T1"), belge("P3", "T9"), belge("P4", "")])
        self.assertEqual(f, [("Yüksek", "Aynı teklife birden çok poliçe")] * 2 + [("Yüksek", "Teklif bulunamadı")] * 2)
        self.assertEqual([o["sonuc"] for o in r["ozet"]], ["Fark var", "Fark var", "Teklif yok", "Teklif yok"])
        self.assertEqual(r["farklar"][0]["policede"], "P1, P2")
        r = main.karsilastir({"T1": belge("T1"), "T2": belge("T2")}, {"P1": belge("P1", "T1")})
        self.assertEqual([t["no"] for t in r["uretilmemis"]], ["T2"])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "p.xlsx"
            r = main.calistir(ORNEK / "teklifler.csv", ORNEK / "policeler.csv", cikti)
            d = {o["p"]["no"]: (o["sonuc"], sorted(f["konu"] for f in r["farklar"] if f["police"] == o["p"]["no"])) for o in r["ozet"]}
            self.assertEqual(d["P-770001"], ("Uyumlu", []))
            self.assertEqual(d["P-770002"][1], ["Muafiyet farklı", "Teminat poliçede yok", "Özel şart poliçede yok"])
            self.assertEqual(d["P-770003"][1], ["Bedel farklı", "Prim farklı", "Teminat teklifte yok"])
            self.assertEqual(d["P-770004"][1], ["Bitiş tarihi farklı", "Sigortalı adı farklı", "Özel şart teklifte yok"])       # 0,50 TL prim farkı tolerans içinde
            self.assertEqual((d["P-770005"][1], d["P-770006"][1], d["P-770007"]), (["Aynı teklife birden çok poliçe"],) * 2 + (("Teklif yok", ["Teklif bulunamadı"]),))
            o3 = next(o for o in r["ozet"] if o["p"]["no"] == "P-770003")
            self.assertEqual((o3["teklif_prim"], o3["police_prim"], len(r["farklar"])), (D(225000), D(220750), 12))
            self.assertEqual(([t["no"] for t in r["uretilmemis"]], r["uyarilar"]), (["TK-3006"], []))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Poliçe Özeti", "Farklar", "Üretilmemiş Teklifler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            t, p = Path(tmp) / "t.csv", Path(tmp) / "p.csv"
            t.write_text("Teklif No;Teminat;Bedel;Prim;Başlangıç;Özel Şartlar\nT1;Yangın;1.000,00;10,00;01.11.2026;Kloz A\nT1;Deprem;bin;20;yarın;Kloz B\nT1;Yangın;5;5;;\n;Cam;1;1;;\n",
                         encoding="utf-8")
            b, uy = main.oku(t, "teklif")
            self.assertEqual((list(b), list(b["T1"]["teminat"]), b["T1"]["teminat"]["deprem"]["bedel"], b["T1"]["bas"], sorted(b["T1"]["sart"])),
                             (["T1"], ["yangin", "deprem"], None, date(2026, 11, 1), ["kloz a", "kloz b"]))
            self.assertEqual(sorted(u["tur"] for u in uy), ["Bedel okunamadı", "Tarih okunamadı", "Teminat tekrar ediyor"])
            p.write_text("Poliçe No;Teklif No;Teminat;Bedel;Prim\nP1;T1;Yangın;1.000,00;10,00\n", encoding="utf-8")
            c = str(Path(tmp) / "x.xlsx")
            self.assertEqual(main.main(["--teklifler", str(t), "--policeler", str(p), "--cikti", c]), 0)
            self.assertEqual(main.main(["--teklifler", str(t), "--policeler", str(p), "--cikti", c, "--tolerans", "-1"]), 2)
            self.assertEqual(main.main(["--policeler", str(Path(tmp) / "yok.csv")]), 1)
            p.write_text("Poliçe No;Teklif No;Teminat\n", encoding="utf-8")
            self.assertEqual(main.main(["--teklifler", str(t), "--policeler", str(p), "--cikti", c]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
