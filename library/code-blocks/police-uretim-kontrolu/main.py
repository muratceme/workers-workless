"""
Poliçe Üretim Kontrolü — Workers / Workless kod bloğu
Teknik (Risk Kabul / Underwriting) › Teknik Uzman Yardımcısı

Üretilen poliçeleri onaylı tekliflerle teminat satırı düzeyinde karşılaştırır.
  Eşleştirme : poliçe → Teklif No;  teminat → teminat adı (büyük-küçük harf ve Türkçe karakter ayrımı yapılmaz)
  Kontroller : eksik / fazla teminat, bedel, prim, muafiyet, sigortalı, başlangıç-bitiş tarihi, özel şartlar (klozlar)
  Tolerans   : bedel ve prim farkı --tolerans TL'yi (varsayılan 1) aşarsa fark sayılır
  Ayrıca     : aynı teklife birden çok poliçe, poliçesi üretilmemiş teklifler
Kod poliçe metnini okumaz; yalnız iki tablodaki alanları karşılaştırır.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"police": ("police no", "police numarasi"), "teklif": ("teklif no", "teklif numarasi"), "sigortali": ("sigortali", "sigortali adi", "musteri"),
         "teminat": ("teminat", "teminat adi"), "bedel": ("bedel", "sigorta bedeli", "teminat bedeli", "limit"), "prim": ("prim", "net prim", "teminat primi"),
         "muafiyet": ("muafiyet",), "bas": ("baslangic", "baslangic tarihi", "police baslangic"), "bit": ("bitis", "bitis tarihi", "police bitis"),
         "sart": ("ozel sartlar", "ozel sart", "klozlar", "kloz")}
Z = Decimal(0)


def sartlar(x) -> dict[str, str]:
    return {katla(s): s.strip() for s in re.split(r"[|;\n]", metin(x)) if katla(s)}


def oku(yol: Path, anahtar: str) -> tuple[dict, list[dict]]:
    """anahtar = 'teklif' veya 'police'. Aynı numaralı satırlar tek belge altında toplanır."""
    uy, belgeler = [], {}
    for r in kayitlar(yol, SUTUN, (anahtar, "teminat")):
        no, tem = metin(r[anahtar]), metin(r["teminat"])
        if not no:
            continue
        b = belgeler.setdefault(no, {"no": no, "teklif": "", "sigortali": "", "bas": None, "bit": None, "sart": {}, "teminat": {}})
        kim = f"{no} (satır {r['_satir']})"
        b["teklif"] = b["teklif"] or metin(r["teklif"])
        b["sigortali"] = b["sigortali"] or metin(r["sigortali"])
        b["sart"] |= sartlar(r["sart"])
        for a, ad in (("bas", "Başlangıç"), ("bit", "Bitiş")):
            t = tarih(r[a])
            if metin(r[a]) and not t:
                uy.append(uyari("Orta", "Tarih okunamadı", kim, f"{yol.name}: {ad} '{metin(r[a])}'."))
            b[a] = b[a] or t
        if not tem:
            continue
        bedel, prim = para(r["bedel"]), para(r["prim"])
        for v, ham, ad in ((bedel, r["bedel"], "Bedel"), (prim, r["prim"], "Prim")):
            if v is None and metin(ham):
                uy.append(uyari("Yüksek", f"{ad} okunamadı", kim, f"{yol.name}: {tem} · '{metin(ham)}'. Karşılaştırmada boş sayıldı."))
        if katla(tem) in b["teminat"]:
            uy.append(uyari("Yüksek", "Teminat tekrar ediyor", kim, f"{yol.name}: '{tem}' aynı belgede iki kez yazılmış; ilk satır kullanıldı."))
            continue
        b["teminat"][katla(tem)] = {"ad": tem, "bedel": bedel, "prim": prim, "muafiyet": metin(r["muafiyet"])}
    return belgeler, uy


def tl(x) -> str:
    return "boş" if x is None else f"{yuvarla(x):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def karsilastir(teklifler: dict, policeler: dict, tolerans: Decimal = Decimal(1)) -> dict:
    farklar, ozet = [], []
    sayac = defaultdict(list)
    for p in policeler.values():
        sayac[p["teklif"]].append(p["no"])

    def ekle(p, duzey, konu, teminat, tk, pl, ac=""):
        farklar.append({"police": p["no"], "teklif": p["teklif"], "duzey": duzey, "konu": konu, "teminat": teminat, "teklifte": tk, "policede": pl, "aciklama": ac})

    for p in policeler.values():
        n0, t = len(farklar), teklifler.get(p["teklif"])
        pp = sum((x["prim"] or Z for x in p["teminat"].values()), Z)
        if not t:
            ekle(p, "Yüksek", "Teklif bulunamadı", "", "", p["teklif"] or "boş", "Poliçedeki teklif numarası onaylı teklif listesinde yok; poliçe karşılaştırılamadı.")
            ozet.append({"p": p, "sonuc": "Teklif yok", "fark": 1, "teklif_prim": None, "police_prim": pp})
            continue
        if len(sayac[p["teklif"]]) > 1:
            ekle(p, "Yüksek", "Aynı teklife birden çok poliçe", "", "", ", ".join(sayac[p["teklif"]]), "Mükerrer üretim olabilir; zeyil veya yenileme değilse birini iptal edin.")
        if katla(t["sigortali"]) != katla(p["sigortali"]) and t["sigortali"] and p["sigortali"]:
            ekle(p, "Orta", "Sigortalı adı farklı", "", t["sigortali"], p["sigortali"])
        for a, ad in (("bas", "Başlangıç tarihi farklı"), ("bit", "Bitiş tarihi farklı")):
            if t[a] and p[a] and t[a] != p[a]:
                ekle(p, "Yüksek", ad, "", f"{t[a]:%d.%m.%Y}", f"{p[a]:%d.%m.%Y}")
        for k, x in t["teminat"].items():
            y = p["teminat"].get(k)
            if not y:
                ekle(p, "Yüksek", "Teminat poliçede yok", x["ad"], f"bedel {tl(x['bedel'])}, prim {tl(x['prim'])}", "", "Onaylı teklifteki teminat poliçeye girilmemiş.")
                continue
            for alan, ad in (("bedel", "Bedel farklı"), ("prim", "Prim farklı")):
                a, b = x[alan], y[alan]
                if (a is None) != (b is None) or (a is not None and abs(a - b) > tolerans):
                    ekle(p, "Yüksek", ad, x["ad"], tl(a), tl(b), f"Fark {tl(b - a)}." if a is not None and b is not None else "")
            if katla(x["muafiyet"]) != katla(y["muafiyet"]):
                ekle(p, "Orta", "Muafiyet farklı", x["ad"], x["muafiyet"] or "boş", y["muafiyet"] or "boş")
        for k, y in p["teminat"].items():
            if k not in t["teminat"]:
                ekle(p, "Yüksek", "Teminat teklifte yok", y["ad"], "", f"bedel {tl(y['bedel'])}, prim {tl(y['prim'])}", "Poliçeye onaylı teklifte olmayan teminat eklenmiş.")
        for k, s in t["sart"].items():
            if k not in p["sart"]:
                ekle(p, "Yüksek", "Özel şart poliçede yok", "", s, "", "Onayda istenen özel şart / kloz poliçeye yazılmamış.")
        for k, s in p["sart"].items():
            if k not in t["sart"]:
                ekle(p, "Orta", "Özel şart teklifte yok", "", "", s, "Poliçeye onayda olmayan özel şart / kloz eklenmiş.")
        tp = sum((x["prim"] or Z for x in t["teminat"].values()), Z)
        ozet.append({"p": p, "sonuc": "Fark var" if len(farklar) > n0 else "Uyumlu", "fark": len(farklar) - n0, "teklif_prim": tp, "police_prim": pp})
    uretilmemis = [t for no, t in teklifler.items() if no not in sayac]
    return {"ozet": ozet, "farklar": farklar, "uretilmemis": uretilmemis}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Poliçe Özeti"
    _baslik(ws, ["Poliçe No", "Teklif No", "Sigortalı", "Sonuç", "Fark Sayısı", "Teklif Primi", "Poliçe Primi", "Prim Farkı", "Kontrol Eden (siz doldurun)"], (16, 14, 32, 13, 9, 16, 16, 14, 18))
    renk = {"Uyumlu": YESIL, "Fark var": KIRMIZI, "Teklif yok": KIRMIZI}
    for o in sorted(r["ozet"], key=lambda o: (o["sonuc"] == "Uyumlu", -o["fark"])):
        fark = o["police_prim"] - o["teklif_prim"] if o["teklif_prim"] is not None else None
        _satir(ws, [o["p"]["no"], o["p"]["teklif"], o["p"]["sigortali"], o["sonuc"], o["fark"], o["teklif_prim"], o["police_prim"], fark, ""], para_s=(6, 7, 8), dolgu={4: renk[o["sonuc"]]})
    ws = wb.create_sheet("Farklar")
    _baslik(ws, ["Poliçe No", "Teklif No", "Düzey", "Konu", "Teminat", "Teklifte", "Poliçede", "Açıklama", "Düzeltildi (siz doldurun)"], (16, 14, 9, 30, 26, 32, 32, 60, 14))
    for f in sorted(r["farklar"], key=lambda f: (f["duzey"] != "Yüksek", f["police"])):
        _satir(ws, [f["police"], f["teklif"], f["duzey"], f["konu"], f["teminat"], f["teklifte"], f["policede"], f["aciklama"], ""], dolgu={3: RENK[f["duzey"]]}, sar=(6, 7, 8))
    ws = wb.create_sheet("Üretilmemiş Teklifler")
    _baslik(ws, ["Teklif No", "Sigortalı", "Başlangıç", "Teminat Sayısı", "Toplam Prim"], (14, 32, 12, 10, 16))
    for t in r["uretilmemis"]:
        _satir(ws, [t["no"], t["sigortali"], t["bas"], len(t["teminat"]), sum((x["prim"] or Z for x in t["teminat"].values()), Z)], para_s=(5,), tarih_s=(3,))
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(teklif_yol: Path, police_yol: Path, cikti: Path, tolerans: Decimal = Decimal(1)) -> dict:
    teklifler, uy1 = oku(teklif_yol, "teklif")
    policeler, uy2 = oku(police_yol, "police")
    if not policeler:
        raise ValueError(f"{police_yol.name}: kontrol edilecek poliçe yok.")
    r = karsilastir(teklifler, policeler, tolerans)
    r["uyarilar"] = sirala(uy1 + uy2)
    rapor_yaz(cikti, r, uy1 + uy2)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Üretilen poliçeleri onaylı tekliflerle teminat, bedel, prim ve özel şart bazında karşılaştırır.")
    p.add_argument("--teklifler", type=Path, default=ORNEK / "teklifler.csv", help="Teklif No, Sigortalı, Teminat, Bedel, Prim, Muafiyet, Başlangıç, Bitiş, Özel Şartlar")
    p.add_argument("--policeler", type=Path, default=ORNEK / "policeler.csv", help="Poliçe No, Teklif No + teklif dosyasıyla aynı sütunlar")
    p.add_argument("--tolerans", default="1", help="Bedel ve prim için kabul edilen fark, TL (varsayılan 1)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "police_uretim_kontrol.xlsx")
    a = p.parse_args(argv)
    for y in (a.teklifler, a.policeler):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    tol = para(a.tolerans)
    if tol is None or tol < 0:
        print("[X] --tolerans sıfır veya pozitif bir sayı olmalı")
        return 2
    try:
        r = calistir(a.teklifler, a.policeler, a.cikti, tol)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(r['ozet'])} poliçe · {sum(o['sonuc'] == 'Uyumlu' for o in r['ozet'])} uyumlu · {sum(o['sonuc'] != 'Uyumlu' for o in r['ozet'])} farklı · "
          f"{len(r['farklar'])} fark · {len(r['uretilmemis'])} teklifin poliçesi yok")
    for o in r["ozet"]:
        if o["sonuc"] != "Uyumlu":
            print(f"[!] {o['p']['no']}: " + "; ".join(dict.fromkeys(f["konu"] for f in r["farklar"] if f["police"] == o["p"]["no"])))
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
