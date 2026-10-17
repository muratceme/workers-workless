"""
Stok Sorgu ve Transfer Talebi — Gerisi Bende kod bloğu
Mağaza Operasyonları › Satış Danışmanı
Mağazada olmayan ürün / renk / bedenin diğer mağaza stoklarını listeler ve transfer talebi formatına çevirir.
  Verilebilir adet = stok − rezerve − kaynak mağazada bırakılacak adet (--birak, varsayılan 1; son ürünü vitrinden çekmemek için)
  Kaynak sırası    = aynı şehir → aynı bölge → verilebilir adedi çok olan → mağaza adı
  Talep birden fazla mağazadan karşılanabilir (--tek-kaynak verilirse yalnız tamamını tek başına karşılayabilen mağaza)
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
STOK = {"magaza": ("magaza", "magaza adi", "lokasyon", "depo"), "kod": ("urun kodu", "stok kodu", "model kodu", "sku", "barkod"), "ad": ("urun adi", "urun", "aciklama"),
        "renk": ("renk", "renk kodu"), "beden": ("beden", "numara", "olcu"), "stok": ("stok", "eldeki stok", "miktar", "adet"), "rezerve": ("rezerve", "ayrilan", "bloke")}
TALEP = {"magaza": ("talep eden magaza", "talep eden", "magaza", "hedef magaza"), "kod": ("urun kodu", "stok kodu", "model kodu", "sku", "barkod"), "renk": ("renk", "renk kodu"),
         "beden": ("beden", "numara", "olcu"), "adet": ("adet", "miktar", "istenen"), "musteri": ("musteri", "not", "aciklama")}
MAGAZA = {"magaza": ("magaza", "magaza adi"), "bolge": ("bolge",), "sehir": ("sehir", "il")}


def _anahtar(kod, renk, beden) -> tuple:
    return (katla(kod), katla(renk), katla(beden))


def stok_oku(yol: Path) -> tuple[dict, list[dict]]:
    uy, d = [], {}
    for r in kayitlar(yol, STOK, ("magaza", "kod", "stok")):
        if not metin(r["magaza"]) or not metin(r["kod"]):
            continue
        s, rz = para(r["stok"]), (para(r["rezerve"]) if metin(r["rezerve"]) else Decimal(0))
        if s is None or rz is None or rz < 0:
            uy.append(uyari("Yüksek", "Stok satırı okunamadı", f"{metin(r['magaza'])} · {metin(r['kod'])}", f"Satır {r['_satir']}: stok ya da rezerve okunamadı."))
            continue
        if s < 0:
            uy.append(uyari("Orta", "Negatif stok", f"{metin(r['magaza'])} · {metin(r['kod'])} {metin(r['renk'])} {metin(r['beden'])}".strip(), f"Stok {s}; sayım / kayıt hatası olabilir, 0 kabul edildi."))
            s = Decimal(0)
        k = _anahtar(r["kod"], r["renk"], r["beden"])
        urun = d.setdefault(k, {"kod": metin(r["kod"]), "ad": metin(r["ad"]), "renk": metin(r["renk"]), "beden": metin(r["beden"]), "magazalar": {}})
        urun["ad"] = urun["ad"] or metin(r["ad"])
        m = urun["magazalar"].setdefault(katla(r["magaza"]), {"magaza": metin(r["magaza"]), "stok": Decimal(0), "rezerve": Decimal(0)})
        m["stok"] += s
        m["rezerve"] += rz
    return d, uy


def talep_oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, TALEP, ("magaza", "kod")):
        if not metin(r["magaza"]) and not metin(r["kod"]):
            continue
        a = para(r["adet"]) if metin(r["adet"]) else Decimal(1)
        if not metin(r["magaza"]) or not metin(r["kod"]) or a is None or a <= 0 or a != int(a):
            uy.append(uyari("Yüksek", "Talep satırı okunamadı", metin(r["kod"]) or f"Satır {r['_satir']}", f"Satır {r['_satir']}: mağaza, ürün kodu ya da adet okunamadı."))
            continue
        l.append({"magaza": metin(r["magaza"]), "kod": metin(r["kod"]), "renk": metin(r["renk"]), "beden": metin(r["beden"]), "adet": int(a), "musteri": metin(r["musteri"]), "satir": r["_satir"]})
    return l, uy


def magaza_oku(yol: Path) -> dict:
    return {katla(r["magaza"]): {"bolge": katla(r["bolge"]), "sehir": katla(r["sehir"])} for r in kayitlar(yol, MAGAZA, ("magaza",)) if metin(r["magaza"])}


def dagit(talepler: list[dict], stok: dict, magazalar: dict | None = None, birak: int = 1, tek_kaynak: bool = False) -> dict:
    if not talepler:
        raise ValueError("Talep yok.")
    magazalar = magazalar or {}
    uy, transfer, sorgu, karsilanmayan = [], [], [], []
    ayrilan = {}
    for t in talepler:
        kim = f"{t['magaza']} · {t['kod']} {t['renk']} {t['beden']}".strip()
        urun = stok.get(_anahtar(t["kod"], t["renk"], t["beden"]))
        if urun is None:
            karsilanmayan.append({**t, "neden": "Ürün / renk / beden stok listesinde yok", "eksik": t["adet"]})
            uy.append(uyari("Orta", "Ürün bulunamadı", kim, "Stok listesinde bu ürün / renk / beden kombinasyonu yok; kodu ve beden yazımını kontrol edin."))
            continue
        hedef = magazalar.get(katla(t["magaza"]), {})
        kendi = urun["magazalar"].get(katla(t["magaza"]))
        if kendi and kendi["stok"] - kendi["rezerve"] >= t["adet"]:
            uy.append(uyari("Bilgi", "Talep eden mağazada stok var", kim, f"Mağazada {kendi['stok'] - kendi['rezerve']} adet kullanılabilir stok görünüyor; reyon / depo kontrolü yapın."))
        adaylar = []
        for km, m in urun["magazalar"].items():
            if km == katla(t["magaza"]):
                continue
            ver = int(m["stok"] - m["rezerve"] - ayrilan.get((km, _anahtar(t["kod"], t["renk"], t["beden"])), 0)) - birak
            bilgi = magazalar.get(km, {})
            adaylar.append({"magaza": m["magaza"], "km": km, "stok": m["stok"], "verilebilir": max(ver, 0),
                            "yakinlik": (0 if hedef.get("sehir") and bilgi.get("sehir") == hedef.get("sehir") else 1, 0 if hedef.get("bolge") and bilgi.get("bolge") == hedef.get("bolge") else 1)})
        adaylar.sort(key=lambda a: (a["yakinlik"], -a["verilebilir"], katla(a["magaza"])))
        sorgu.append({**t, "ad": urun["ad"], "stoklar": [(a["magaza"], a["stok"], a["verilebilir"]) for a in adaylar if a["stok"] > 0]})
        kalan = t["adet"]
        uygun = [a for a in adaylar if a["verilebilir"] > 0]
        stok_var = bool(uygun)
        if tek_kaynak:
            uygun = [a for a in uygun if a["verilebilir"] >= kalan][:1]
        for a in uygun:
            if kalan <= 0:
                break
            al = min(kalan, a["verilebilir"])
            ayrilan[(a["km"], _anahtar(t["kod"], t["renk"], t["beden"]))] = ayrilan.get((a["km"], _anahtar(t["kod"], t["renk"], t["beden"])), 0) + al
            transfer.append({"kaynak": a["magaza"], "hedef": t["magaza"], "kod": urun["kod"], "ad": urun["ad"], "renk": urun["renk"], "beden": urun["beden"], "adet": al, "musteri": t["musteri"],
                             "yakin": "Aynı şehir" if a["yakinlik"][0] == 0 else ("Aynı bölge" if a["yakinlik"][1] == 0 else "")})
            kalan -= al
        if kalan > 0:
            neden = "Diğer mağazalarda verilebilir stok yok" if not stok_var else ("Tek mağazada yeterli stok yok" if tek_kaynak else "Kısmen karşılandı")
            karsilanmayan.append({**t, "neden": neden, "eksik": kalan})
            uy.append(uyari("Orta", "Talep karşılanamadı", kim, f"{t['adet']} adetten {kalan} adet karşılanamadı ({neden}); merkez depo ya da tedarik talebi gerekebilir."))
    return {"transfer": transfer, "sorgu": sorgu, "karsilanmayan": karsilanmayan, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Transfer Talepleri"
    _baslik(ws, ["Kaynak Mağaza", "Hedef Mağaza", "Ürün Kodu", "Ürün Adı", "Renk", "Beden", "Adet", "Yakınlık", "Müşteri / Not"], (20, 20, 14, 28, 10, 7, 6, 11, 24))
    for t in sorted(r["transfer"], key=lambda t: (katla(t["kaynak"]), katla(t["hedef"]), t["kod"])):
        _satir(ws, [t["kaynak"], t["hedef"], t["kod"], t["ad"], t["renk"], t["beden"], t["adet"], t["yakin"], t["musteri"]])
    ws = wb.create_sheet("Stok Sorgusu")
    _baslik(ws, ["Talep Eden", "Ürün Kodu", "Ürün Adı", "Renk", "Beden", "İstenen", "Stoku Olan Mağazalar (stok / verilebilir)"], (20, 14, 28, 10, 7, 7, 80))
    for s in r["sorgu"]:
        _satir(ws, [s["magaza"], s["kod"], s["ad"], s["renk"], s["beden"], s["adet"], "; ".join(f"{m} ({st} / {v})" for m, st, v in s["stoklar"]) or "Stok yok"], sar=(7,))
    ws = wb.create_sheet("Karşılanamayan")
    _baslik(ws, ["Talep Eden", "Ürün Kodu", "Renk", "Beden", "İstenen", "Eksik", "Neden"], (20, 14, 10, 7, 7, 6, 40))
    for k in r["karsilanmayan"]:
        _satir(ws, [k["magaza"], k["kod"], k["renk"], k["beden"], k["adet"], k["eksik"], k["neden"]])
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(stok_yol: Path, talep_yol: Path, cikti: Path, magaza_yol: Path | None = None, birak: int = 1, tek_kaynak: bool = False) -> dict:
    st, uy1 = stok_oku(stok_yol)
    tl, uy2 = talep_oku(talep_yol)
    r = dagit(tl, st, magaza_oku(magaza_yol) if magaza_yol else None, birak, tek_kaynak)
    r["uyarilar"] = sirala(uy1 + uy2 + r["uyarilar"])
    rapor_yaz(cikti, r)
    with open(cikti.with_name(cikti.stem + "_transfer.csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Kaynak Mağaza", "Hedef Mağaza", "Ürün Kodu", "Renk", "Beden", "Adet", "Not"])
        for t in r["transfer"]:
            w.writerow([t["kaynak"], t["hedef"], t["kod"], t["renk"], t["beden"], t["adet"], t["musteri"]])
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Mağazada olmayan beden / rengin diğer mağaza stoklarını bulur, transfer talebi hazırlar.")
    p.add_argument("--stok", type=Path, help="Mağaza, Ürün Kodu, Ürün Adı, Renk, Beden, Stok, Rezerve")
    p.add_argument("--talepler", type=Path, help="Talep Eden Mağaza, Ürün Kodu, Renk, Beden, Adet, Müşteri / Not")
    p.add_argument("--magazalar", type=Path, help="Mağaza, Bölge, Şehir (isteğe bağlı; yakın mağaza önceliği için)")
    p.add_argument("--birak", type=int, default=1, help="Kaynak mağazada bırakılacak en az adet (varsayılan 1)")
    p.add_argument("--tek-kaynak", action="store_true", help="Talebi yalnız tamamını karşılayabilen tek mağazadan iste")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "transfer_talebi.xlsx")
    a = p.parse_args(argv)
    ornek = a.stok is None
    sy, ty = a.stok or ORNEK / "stok.csv", a.talepler or ORNEK / "talepler.csv"
    my = a.magazalar or (ORNEK / "magazalar.csv" if ornek else None)
    if a.birak < 0:
        print("[X] --birak 0 ya da üstü olmalı")
        return 2
    for y in (sy, ty, my):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(sy, ty, a.cikti, my, a.birak, a.tek_kaynak)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(r['sorgu']) + sum(1 for k in r['karsilanmayan'] if k['neden'].startswith('Ürün'))} talep · {len(r['transfer'])} transfer satırı ({sum(t['adet'] for t in r['transfer'])} adet) · "
          f"karşılanamayan {len(r['karsilanmayan'])}")
    for t in r["transfer"][:8]:
        print(f"     {t['kaynak']} → {t['hedef']}: {t['kod']} {t['renk']} {t['beden']} × {t['adet']}")
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Talep: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
