import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def stok(**m):
    return {("k1", "siyah", "m"): {"kod": "K1", "ad": "Ürün", "renk": "Siyah", "beden": "M",
                                   "magazalar": {main.katla(a): {"magaza": a, "stok": D(s), "rezerve": D(r)} for a, (s, r) in m.items()}}}


def t(magaza, adet=1, kod="K1"):
    return {"magaza": magaza, "kod": kod, "renk": "siyah", "beden": "M", "adet": adet, "musteri": "", "satir": 2}


MAG = {"a": {"bolge": "b1", "sehir": "s1"}, "b": {"bolge": "b1", "sehir": "s2"}, "c": {"bolge": "b2", "sehir": "s1"}, "d": {"bolge": "b9", "sehir": "s9"}}


class DagitimTesti(unittest.TestCase):
    def test_oncelik(self):
        st = stok(A=(1, 0), B=(3, 0), C=(2, 0), D=(9, 0))
        r = main.dagit([t("A", 2)], st, MAG)
        # A s1/b1: önce aynı şehir C (2-1=1), sonra aynı bölge B (3-1=2)
        self.assertEqual([(x["kaynak"], x["adet"], x["yakin"]) for x in r["transfer"]], [("C", 1, "Aynı şehir"), ("B", 1, "Aynı bölge")])
        r = main.dagit([t("A", 2)], st)
        self.assertEqual([(x["kaynak"], x["adet"]) for x in r["transfer"]], [("D", 2)])
        r = main.dagit([t("A", 3)], st, MAG, tek_kaynak=True)
        self.assertEqual([(x["kaynak"], x["adet"]) for x in r["transfer"]], [("D", 3)])
        r = main.dagit([t("A", 9)], st, MAG, tek_kaynak=True)
        self.assertEqual((r["transfer"], r["karsilanmayan"][0]["neden"]), ([], "Tek mağazada yeterli stok yok"))

    def test_birak_rezerve_ve_ardisik(self):
        st = stok(A=(0, 0), B=(3, 1))
        r = main.dagit([t("A", 1), t("C", 1), t("A", 1, kod="X")], st, birak=1)
        self.assertEqual([(x["kaynak"], x["hedef"], x["adet"]) for x in r["transfer"]], [("B", "A", 1)])
        self.assertEqual([(k["magaza"], k["eksik"], k["neden"][:5]) for k in r["karsilanmayan"]], [("C", 1, "Diğer"), ("A", 1, "Ürün ")])
        r = main.dagit([t("A", 5)], st, birak=0)
        self.assertEqual((r["transfer"][0]["adet"], r["karsilanmayan"][0]["neden"], r["karsilanmayan"][0]["eksik"]), (2, "Kısmen karşılandı", 3))
        r = main.dagit([t("B", 1)], st)
        self.assertEqual([u["tur"] for u in r["uyarilar"]], ["Talep eden mağazada stok var", "Talep karşılanamadı"])
        with self.assertRaises(ValueError):
            main.dagit([], st)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "stok.csv", ORNEK / "talepler.csv", cikti, ORNEK / "magazalar.csv")
            self.assertEqual((len(r["transfer"]), sum(x["adet"] for x in r["transfer"]), len(r["karsilanmayan"])), (6, 8, 3))
            self.assertEqual(r["transfer"][0]["kaynak"], "Bakırköy AVM")
            self.assertEqual(load_workbook(cikti).sheetnames, ["Transfer Talepleri", "Stok Sorgusu", "Karşılanamayan", "Uyarılar"])
            self.assertEqual(len((Path(tmp) / "t_transfer.csv").read_text(encoding="utf-8-sig").splitlines()), 7)

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            sy, ty = Path(tmp) / "s.csv", Path(tmp) / "t.csv"
            sy.write_text("Mağaza;SKU;Renk;Beden;Stok\nA;K1;Siyah;M;2\nA;K1;siyah;m;1\nB;K1;Siyah;M;-1\nC;K1;Siyah;M;x\n", encoding="utf-8")
            ty.write_text("Talep Eden;SKU;Renk;Beden;Adet\nD;K1;Siyah;M;\nD;K1;Siyah;M;1,5\n;;;;\n", encoding="utf-8")
            s, uy = main.stok_oku(sy)
            self.assertEqual((s[("k1", "siyah", "m")]["magazalar"]["a"]["stok"], s[("k1", "siyah", "m")]["magazalar"]["b"]["stok"], [u["tur"] for u in uy]),
                             (D(3), D(0), ["Negatif stok", "Stok satırı okunamadı"]))
            l, uy = main.talep_oku(ty)
            self.assertEqual(([x["adet"] for x in l], len(uy)), ([1], 1))
            c = str(Path(tmp) / "x.xlsx")
            self.assertEqual(main.main(["--stok", str(sy), "--talepler", str(ty), "--birak", "0", "--tek-kaynak", "--cikti", c]), 0)
            self.assertEqual(main.main(["--birak", "-1", "--cikti", c]), 2)
            self.assertEqual(main.main(["--stok", str(Path(tmp) / "yok.csv")]), 1)
            ty.write_text("Talep Eden;SKU\n;\n", encoding="utf-8")
            self.assertEqual(main.main(["--stok", str(sy), "--talepler", str(ty), "--cikti", c]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
