"""
Kesinti ve İtiraz Takibi — Workers / Workless kod bloğu
Anlaşmalı Kurumlar ve Faturalama › Anlaşmalı Kurumlar Uzmanı

Sigorta şirketlerinin yaptığı kesintileri neden bazında sınıflandırır, itiraz edilebilir olanları ve süreleri takip eder.
  Neden kategorisi   = kural dosyasındaki ilk eşleşen satır (anahtar kelimelerden biri kesinti nedeni metninde geçiyorsa); eşleşme yoksa "Diğer"
  İtiraz edilebilir  = kural dosyasındaki Evet / Hayır; "Diğer" kategorisi "İncelenecek" sayılır
  Son itiraz günü    = kesinti tarihi + şirketin itiraz süresi (süre dosyası; şirket orada yoksa --itiraz-gun; ikisi de yoksa süre takibi yapılmaz)
  Durum: Kabul edildi · Kısmi kabul · Reddedildi (itiraz sonucu yazılmış) · Sonuç bekleniyor (itiraz tarihi var, sonuç yok) · İtiraz edilmeyecek (kategori "Hayır")
         Süre geçti (itiraz yok, rapor tarihi > son gün) · Süre yaklaşıyor (kalan gün ≤ --yaklasan-gun 5) · İtiraz bekliyor
  Geri kazanım %     = geri alınan tutar ÷ sonuçlanan itirazların kesinti tutarı × 100
İtiraz süresi ve hangi kesintiye itiraz edilebileceği şirketle yaptığınız sözleşmeye bağlıdır; kod bir süre veya hukuki değerlendirme varsaymaz. Hasta bilgisi gerekmez.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"sirket": ("sigorta sirketi", "sirket", "kurum", "sigorta"), "fatura": ("fatura no", "fatura", "fatura numarasi", "dosya no"), "tarih": ("kesinti tarihi", "bildirim tarihi", "tarih"),
         "tutar": ("kesinti tutari", "kesinti", "tutar"), "neden": ("kesinti nedeni", "neden", "aciklama", "kesinti aciklamasi"), "itiraz": ("itiraz tarihi",),
         "sonuc": ("itiraz sonucu", "sonuc", "itiraz durumu"), "geri": ("geri alinan tutar", "geri alinan", "iade edilen tutar", "odenen tutar")}
KURAL = {"kategori": ("neden kategorisi", "kategori"), "kelime": ("anahtar kelimeler", "anahtar kelime", "kelimeler"), "edilebilir": ("itiraz edilebilir", "itiraz", "itiraz edilebilir mi")}
SURE = {"sirket": SUTUN["sirket"], "sure": ("itiraz suresi", "itiraz suresi gun", "sure", "gun")}
SIRA = ("Süre geçti", "Süre yaklaşıyor", "İtiraz bekliyor", "Sonuç bekleniyor", "Kısmi kabul", "Reddedildi", "Kabul edildi", "İtiraz edilmeyecek")
ACIK = ("Süre yaklaşıyor", "İtiraz bekliyor")
Z = Decimal(0)


def kural_oku(yol: Path) -> list[dict]:
    kurallar = []
    for r in kayitlar(yol, KURAL, ("kategori", "kelime")):
        k, kelimeler = metin(r["kategori"]), [katla(x) for x in re.split(r"[,|]", metin(r["kelime"])) if katla(x)]
        if k and kelimeler:
            kurallar.append({"kategori": k, "kelimeler": kelimeler, "edilebilir": "Hayır" if katla(r["edilebilir"]) in ("h", "hayir", "yok", "0") else "Evet"})
    if not kurallar:
        raise ValueError(f"{yol.name}: kural dosyasında kategori ve anahtar kelime yok.")
    return kurallar


def sonuc_coz(x) -> str | None:
    k = katla(x)
    if not k:
        return None
    if "kismi" in k or "kismen" in k:
        return "Kısmi kabul"
    if re.search(r"\b(ret|red|reddedildi|olumsuz)\b", k) or "edilmedi" in k:
        return "Reddedildi"
    if "kabul" in k or "olumlu" in k:
        return "Kabul edildi"
    return ""


def oku(yol: Path, sure_yol: Path | None = None) -> tuple[list[dict], dict, list[dict]]:
    uy, l, sure = [], [], {}
    for r in kayitlar(yol, SUTUN, ("sirket", "tarih", "tutar")):
        s = metin(r["sirket"])
        if not s:
            continue
        kim = f"{s} · {metin(r['fatura']) or 'fatura no yok'} (satır {r['_satir']})"
        t, tutar = tarih(r["tarih"]), para(r["tutar"])
        if t is None or tutar is None or tutar <= 0:
            uy.append(uyari("Yüksek", "Kesinti tarihi veya tutarı okunamadı", kim, f"Tarih '{metin(r['tarih'])}', tutar '{metin(r['tutar'])}'. Kesinti alınmadı."))
            continue
        it, sonuc, geri = tarih(r["itiraz"]), sonuc_coz(r["sonuc"]), para(r["geri"])
        if it is None and metin(r["itiraz"]):
            uy.append(uyari("Yüksek", "İtiraz tarihi okunamadı", kim, f"'{metin(r['itiraz'])}'. İtiraz yapılmamış sayıldı."))
        if sonuc == "":
            uy.append(uyari("Orta", "İtiraz sonucu tanınmadı", kim, f"'{metin(r['sonuc'])}'. Tanınan sonuçlar: Kabul, Kısmi kabul, Ret. Sonuç bekleniyor sayıldı."))
            sonuc = None
        if sonuc and it is None:
            uy.append(uyari("Orta", "Sonuç var, itiraz tarihi yok", kim, "İtiraz sonucu yazılmış ama itiraz tarihi boş."))
        if sonuc == "Kabul edildi" and geri is None:
            geri = tutar
        elif sonuc == "Kısmi kabul" and geri is None:
            uy.append(uyari("Orta", "Kısmi kabulde geri alınan tutar yok", kim, "Geri alınan tutar 0 sayıldı; geri kazanım oranı eksik çıkar."))
        if geri is not None and geri > tutar:
            uy.append(uyari("Orta", "Geri alınan tutar kesintiden büyük", kim, f"Geri alınan {geri:f} > kesinti {tutar:f}."))
        l.append({"sirket": s, "fatura": metin(r["fatura"]), "tarih": t, "tutar": tutar, "neden": metin(r["neden"]), "itiraz": it, "sonuc": sonuc, "geri": (geri or Z) if sonuc else Z})
    if sure_yol:
        for r in kayitlar(sure_yol, SURE, ("sirket", "sure")):
            s, g = metin(r["sirket"]), para(r["sure"])
            if not s:
                continue
            if g is None or g <= 0 or g != int(g):
                uy.append(uyari("Yüksek", "İtiraz süresi okunamadı", s, f"{sure_yol.name} satır {r['_satir']}: '{metin(r['sure'])}'."))
                continue
            sure[katla(s)] = int(g)
    return l, sure, uy


def takip(kesintiler: list[dict], kurallar: list[dict], rapor: date, sure: dict | None = None, itiraz_gun: int | None = None, yaklasan_gun: int = 5) -> dict:
    if not kesintiler:
        raise ValueError("Takip edilecek kesinti yok.")
    sure, edilebilir = sure or {}, {k["kategori"]: k["edilebilir"] for k in kurallar} | {"Diğer": "İncelenecek"}
    for k in kesintiler:
        n = katla(k["neden"])
        k["kategori"] = next((kr["kategori"] for kr in kurallar if any(kel in n for kel in kr["kelimeler"])), "Diğer")
        k["edilebilir"] = edilebilir[k["kategori"]]
        gun = sure.get(katla(k["sirket"]), itiraz_gun)
        k["son_gun"] = k["tarih"] + timedelta(days=gun) if gun is not None else None
        k["kalan"] = (k["son_gun"] - rapor).days if k["son_gun"] and not k["itiraz"] else None
        k["gec_itiraz"] = bool(k["itiraz"] and k["son_gun"] and k["itiraz"] > k["son_gun"])
        if k["sonuc"]:
            k["durum"] = k["sonuc"]
        elif k["itiraz"]:
            k["durum"] = "Sonuç bekleniyor"
        elif k["edilebilir"] == "Hayır":
            k["durum"] = "İtiraz edilmeyecek"
        elif k["kalan"] is not None and k["kalan"] < 0:
            k["durum"] = "Süre geçti"
        else:
            k["durum"] = "Süre yaklaşıyor" if k["kalan"] is not None and k["kalan"] <= yaklasan_gun else "İtiraz bekliyor"

    def topla(liste: list[dict]) -> dict:
        sonuclanan = sum((k["tutar"] for k in liste if k["sonuc"]), Z)
        geri = sum((k["geri"] for k in liste), Z)
        return {"adet": len(liste), "tutar": sum((k["tutar"] for k in liste), Z), "itiraz_edilen": sum((k["tutar"] for k in liste if k["itiraz"] or k["sonuc"]), Z), "sonuclanan": sonuclanan, "geri": geri,
                "geri_kazanim": yuvarla(geri / sonuclanan * 100) if sonuclanan else None, "bekleyen": sum((k["tutar"] for k in liste if k["durum"] in ACIK), Z),
                "sure_gecti": sum((k["tutar"] for k in liste if k["durum"] == "Süre geçti"), Z)}

    genel = topla(kesintiler)
    sirket, neden = defaultdict(list), defaultdict(list)
    for k in kesintiler:
        sirket[katla(k["sirket"])].append(k)
        neden[k["kategori"]].append(k)
    return {"kesintiler": sorted(kesintiler, key=lambda k: (SIRA.index(k["durum"]), k["kalan"] if k["kalan"] is not None else 10**6, -k["tutar"])),
            "sirket": sorted(({"sirket": v[0]["sirket"]} | topla(v) for v in sirket.values()), key=lambda x: -x["tutar"]),
            "neden": sorted(({"kategori": a, "edilebilir": edilebilir[a], "pay": yuvarla(sum((k["tutar"] for k in v), Z) / genel["tutar"] * 100)} | topla(v) for a, v in neden.items()), key=lambda x: -x["tutar"]),
            "genel": genel, "rapor": rapor, "say": {d: sum(k["durum"] == d for k in kesintiler) for d in SIRA}, "ayar": (itiraz_gun, yaklasan_gun),
            "suresiz": sorted({k["sirket"] for k in kesintiler if k["son_gun"] is None})}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    renk = {"Süre geçti": KIRMIZI, "Süre yaklaşıyor": SARI, "İtiraz bekliyor": MAVI, "Kabul edildi": YESIL, "Kısmi kabul": YESIL}
    ws = wb.active
    ws.title = "Kesinti Takibi"
    _baslik(ws, ["Durum", "Sigorta Şirketi", "Fatura No", "Kesinti Tarihi", "Kesinti Tutarı", "Neden Kategorisi", "İtiraz Edilebilir", "Son İtiraz Günü", "Kalan Gün", "İtiraz Tarihi", "Geri Alınan",
                 "Kesinti Nedeni", "Not (siz doldurun)"], (17, 22, 14, 12, 14, 24, 11, 12, 8, 12, 14, 44, 28))
    for k in r["kesintiler"]:
        _satir(ws, [k["durum"], k["sirket"], k["fatura"], k["tarih"], k["tutar"], k["kategori"], k["edilebilir"], k["son_gun"], k["kalan"], k["itiraz"], k["geri"] if k["sonuc"] else None, k["neden"], ""],
               para_s=(5, 11), tarih_s=(4, 8, 10), dolgu={1: renk.get(k["durum"])}, sar=(12,))
    ig, yg = r["ayar"]
    ws.append([])
    ws.append([f"Rapor tarihi {r['rapor']:%d.%m.%Y}. Son itiraz günü = kesinti tarihi + şirketin itiraz süresi" + (f" (dosyada olmayan şirketler için {ig} gün)" if ig is not None else "")
               + f". Kalan gün ≤ {yg} ise 'Süre yaklaşıyor'. Süreler ve itiraz edilebilirlik sözleşmenize göre sizin verdiğiniz bilgilerdir."])
    bas = ["Kesinti Adedi", "Kesinti Tutarı", "İtiraz Edilen", "Sonuçlanan", "Geri Alınan", "Geri Kazanım %", "İtiraz Bekleyen", "Süresi Geçen"]
    deger = lambda x: [x["adet"], x["tutar"], x["itiraz_edilen"], x["sonuclanan"], x["geri"], x["geri_kazanim"], x["bekleyen"], x["sure_gecti"]]   # noqa: E731
    ws = wb.create_sheet("Şirket Özeti")
    _baslik(ws, ["Sigorta Şirketi", *bas], (24, 9, 15, 15, 15, 15, 11, 15, 15))
    for s in r["sirket"]:
        _satir(ws, [s["sirket"], *deger(s)], para_s=(3, 4, 5, 6, 7, 8, 9), dolgu={9: KIRMIZI if s["sure_gecti"] else None})
    _satir(ws, ["Toplam", *deger(r["genel"])], para_s=(3, 4, 5, 6, 7, 8, 9))
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws = wb.create_sheet("Neden Özeti")
    _baslik(ws, ["Neden Kategorisi", "İtiraz Edilebilir", "Tutar Payı %", *bas], (28, 11, 10, 9, 15, 15, 15, 15, 11, 15, 15))
    for n in r["neden"]:
        _satir(ws, [n["kategori"], n["edilebilir"], n["pay"], *deger(n)], para_s=(3, 5, 6, 7, 8, 9, 10, 11))
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, kural_yol: Path, cikti: Path, rapor: date, sure_yol: Path | None = None, itiraz_gun: int | None = None, yaklasan_gun: int = 5) -> dict:
    kurallar = kural_oku(kural_yol)
    kesintiler, sure, uy = oku(yol, sure_yol)
    r = takip(kesintiler, kurallar, rapor, sure, itiraz_gun, yaklasan_gun)
    for s in r["suresiz"]:
        uy.append(uyari("Orta", "İtiraz süresi tanımlı değil", s, "Bu şirket için süre verilmedi; son itiraz günü hesaplanmadı. Sözleşmedeki süreyi süre dosyasına yazın veya --itiraz-gun verin."))
    for k in r["kesintiler"]:
        if k["gec_itiraz"]:
            uy.append(uyari("Orta", "İtiraz süresinden sonra yapılmış", f"{k['sirket']} · {k['fatura']}", f"İtiraz tarihi {k['itiraz']:%d.%m.%Y}, son itiraz günü {k['son_gun']:%d.%m.%Y}."))
        if k["itiraz"] and k["itiraz"] < k["tarih"]:
            uy.append(uyari("Orta", "İtiraz tarihi kesinti tarihinden önce", f"{k['sirket']} · {k['fatura']}", f"İtiraz {k['itiraz']:%d.%m.%Y}, kesinti {k['tarih']:%d.%m.%Y}."))
    n = sum(k["kategori"] == "Diğer" for k in r["kesintiler"])
    if n:
        uy.append(uyari("Bilgi", "Sınıflandırılamayan kesinti nedeni", kural_yol.name, f"{n} kesinti hiçbir anahtar kelimeyle eşleşmedi, 'Diğer' sayıldı. Nedenlere bakıp kural dosyasına kelime ekleyin."))
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Sigorta şirketi kesintilerini sınıflandırır, itiraz sürelerini takip eder.")
    p.add_argument("--kesintiler", type=Path, default=ORNEK / "kesintiler.csv", help="Sigorta Şirketi, Fatura No, Kesinti Tarihi, Kesinti Tutarı, Kesinti Nedeni, İtiraz Tarihi, İtiraz Sonucu, Geri Alınan Tutar")
    p.add_argument("--kurallar", type=Path, default=ORNEK / "neden_kurallari.csv", help="Neden Kategorisi, Anahtar Kelimeler (virgülle), İtiraz Edilebilir (Evet / Hayır)")
    p.add_argument("--sureler", type=Path, help="İsteğe bağlı: Sigorta Şirketi, İtiraz Süresi (gün)")
    p.add_argument("--itiraz-gun", help="Süre dosyasında olmayan şirketler için itiraz süresi (gün); verilmezse bu şirketlerde süre takibi yapılmaz")
    p.add_argument("--yaklasan-gun", default="5", help="Kalan gün bu sayıya eşit veya altındaysa 'Süre yaklaşıyor' (varsayılan 5)")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 12.10.2026)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kesinti_itiraz_takibi.xlsx")
    a = p.parse_args(argv)
    ornek = a.kesintiler == ORNEK / "kesintiler.csv"
    sure_yol = a.sureler or (ORNEK / "itiraz_sureleri.csv" if ornek else None)
    rapor = tarih(a.tarih) if a.tarih else (date(2026, 10, 12) if ornek else date.today())
    ig, yg = para(a.itiraz_gun) if a.itiraz_gun else None, para(a.yaklasan_gun)
    if rapor is None or (a.itiraz_gun and (ig is None or ig <= 0 or ig != int(ig))) or yg is None or yg < 0 or yg != int(yg):
        print("[X] --tarih GG.AA.YYYY biçiminde; --itiraz-gun pozitif, --yaklasan-gun sıfır veya pozitif tam sayı olmalı")
        return 2
    for y in (a.kesintiler, a.kurallar, sure_yol):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.kesintiler, a.kurallar, a.cikti, rapor, sure_yol, None if ig is None else int(ig), int(yg))
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    g, s = r["genel"], r["say"]
    print(f"[OK] {g['adet']} kesinti · {g['tutar']:,.2f} · itiraz edilen {g['itiraz_edilen']:,.2f} · geri alınan {g['geri']:,.2f}" + (f" (geri kazanım %{g['geri_kazanim']})" if g["geri_kazanim"] is not None else ""))
    print(f"[{'!' if s['Süre geçti'] or s['Süre yaklaşıyor'] else 'OK'}] {s['Süre geçti']} süresi geçmiş ({g['sure_gecti']:,.2f}) · {s['Süre yaklaşıyor']} süresi yaklaşan · {s['İtiraz bekliyor']} itiraz bekleyen")
    for k in r["kesintiler"]:
        if k["durum"] == "Süre yaklaşıyor":
            print(f"[!] {k['sirket']} · {k['fatura']}: {k['tutar']:,.2f} · son itiraz günü {k['son_gun']:%d.%m.%Y} ({k['kalan']} gün)")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
