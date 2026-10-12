import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
RAPOR = date(2026, 10, 12)
KURAL = [{"kategori": "Fiyat", "kelimeler": ["fiyat", "tarife"], "edilebilir": "Evet"}, {"kategori": "Provizyon", "kelimeler": ["provizyon"], "edilebilir": "Evet"},
         {"kategori": "Kapsam", "kelimeler": ["kapsam", "limit"], "edilebilir": "Hayır"}]


def ks(neden="Fiyat farkı", gun=-10, tutar=1000, itiraz=None, sonuc=None, geri=0, sirket="A", fatura="F"):
    t = lambda g: None if g is None else RAPOR + timedelta(days=g)   # noqa: E731
    return {"sirket": sirket, "fatura": fatura, "tarih": t(gun), "tutar": D(tutar), "neden": neden, "itiraz": t(itiraz), "sonuc": sonuc, "geri": D(geri)}


def tek(itiraz_gun=30, **kw):
    return main.takip([ks(**kw)], KURAL, RAPOR, itiraz_gun=itiraz_gun)["kesintiler"][0]


class TakipTesti(unittest.TestCase):
    def test_siniflandirma(self):
        self.assertEqual([(x["kategori"], x["edilebilir"]) for x in (tek(neden=n) for n in ("TARİFE üzerinde", "Provizyon limiti aşıldı", "Yıllık limit doldu", "Kaşe yok", ""))],
                         [("Fiyat", "Evet"), ("Provizyon", "Evet"), ("Kapsam", "Hayır"), ("Diğer", "İncelenecek"), ("Diğer", "İncelenecek")])       # ilk eşleşen kural geçerli
        self.assertEqual([main.sonuc_coz(x) for x in ("Kabul", "KISMİ KABUL", "Ret", "Reddedildi", "Kabul edilmedi", "olumlu", "", "inceleniyor")],
                         ["Kabul edildi", "Kısmi kabul", "Reddedildi", "Reddedildi", "Reddedildi", "Kabul edildi", None, ""])

    def test_durum_ve_sure(self):
        self.assertEqual([(x["durum"], x["kalan"]) for x in (tek(gun=g) for g in (-31, -30, -25, -24))], [("Süre geçti", -1), ("Süre yaklaşıyor", 0), ("Süre yaklaşıyor", 5), ("İtiraz bekliyor", 6)])
        x = tek(None, gun=-300)
        self.assertEqual((x["durum"], x["son_gun"], x["kalan"]), ("İtiraz bekliyor", None, None))
        self.assertEqual(main.takip([ks(gun=-20)], KURAL, RAPOR, {"a": 15}, 30)["kesintiler"][0]["durum"], "Süre geçti")               # şirketin süresi varsayılanın önüne geçer
        self.assertEqual((tek(neden="kapsam dışı", gun=-90)["durum"], tek(neden="kaşe yok", gun=-31)["durum"]), ("İtiraz edilmeyecek", "Süre geçti"))
        self.assertEqual([tek(itiraz=-5, sonuc=s)["durum"] for s in (None, "Kabul edildi", "Kısmi kabul", "Reddedildi")], ["Sonuç bekleniyor", "Kabul edildi", "Kısmi kabul", "Reddedildi"])
        self.assertEqual((tek(neden="kapsam dışı", itiraz=-1)["durum"], tek(gun=-40, itiraz=-5)["gec_itiraz"], tek(gun=-40, itiraz=-10)["gec_itiraz"]), ("Sonuç bekleniyor", True, False))

    def test_ozetler(self):
        l = [ks(tutar=4000, gun=-11), ks("Provizyon yok", -27, 2500, -22, "Kabul edildi", 2500), ks("Kapsam dışı", -37, 1200), ks("Fiyat", -41, 700, -32, "Reddedildi"),
             ks("Tarife", -12, 3000, sirket="B"), ks("Provizyon", -32, 5000, -24, "Kısmi kabul", 2000, "B"), ks("Rapor yok", -22, 6000, sirket="B")]
        r = main.takip(l, KURAL, RAPOR, {"a": 30, "b": 15})
        self.assertEqual([(k["tutar"], k["durum"]) for k in r["kesintiler"]],
                         [(D(6000), "Süre geçti"), (D(3000), "Süre yaklaşıyor"), (D(4000), "İtiraz bekliyor"), (D(5000), "Kısmi kabul"), (D(700), "Reddedildi"), (D(2500), "Kabul edildi"), (D(1200), "İtiraz edilmeyecek")])
        self.assertEqual([(s["sirket"], s["adet"], s["tutar"], s["itiraz_edilen"], s["geri"], s["geri_kazanim"], s["bekleyen"], s["sure_gecti"]) for s in r["sirket"]],
                         [("B", 3, D(14000), D(5000), D(2000), D("40.00"), D(3000), D(6000)), ("A", 4, D(8400), D(3200), D(2500), D("78.13"), D(4000), D(0))])
        self.assertEqual([(n["kategori"], n["adet"], n["tutar"], n["pay"]) for n in r["neden"]], [("Fiyat", 3, D(7700), D("34.38")), ("Provizyon", 2, D(7500), D("33.48")), ("Diğer", 1, D(6000), D("26.79")), ("Kapsam", 1, D(1200), D("5.36"))])
        g = r["genel"]
        self.assertEqual((g["tutar"], g["sonuclanan"], g["geri"], g["geri_kazanim"], g["bekleyen"], r["suresiz"]), (D(22400), D(8200), D(4500), D("54.88"), D(7000), []))
        with self.assertRaises(ValueError):
            main.takip([], KURAL, RAPOR)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "k.xlsx"
            r = main.calistir(ORNEK / "kesintiler.csv", ORNEK / "neden_kurallari.csv", cikti, RAPOR, ORNEK / "itiraz_sureleri.csv")
            self.assertEqual({k["fatura"]: (k["durum"], k["kategori"], k["kalan"]) for k in r["kesintiler"]},
                             {"F-101": ("İtiraz bekliyor", "Anlaşma fiyatı farkı", 19), "F-102": ("Kabul edildi", "Provizyon eksik", None), "F-103": ("İtiraz edilmeyecek", "Poliçe kapsamı dışı", -7),
                              "F-201": ("Süre geçti", "Eksik belge", -7), "F-202": ("Süre yaklaşıyor", "Anlaşma fiyatı farkı", 3), "F-203": ("Kısmi kabul", "Provizyon eksik", None),
                              "F-301": ("Sonuç bekleniyor", "Mükerrer fatura", None), "F-302": ("İtiraz bekliyor", "Diğer", None), "F-104": ("Reddedildi", "Anlaşma fiyatı farkı", None)})
            g = r["genel"]
            self.assertEqual((g["tutar"], g["itiraz_edilen"], g["sonuclanan"], g["geri"], g["geri_kazanim"], g["bekleyen"], g["sure_gecti"]), (D(25100), D(10000), D(8200), D(4500), D("54.88"), D(7900), D(6000)))
            self.assertEqual([(n["kategori"], n["tutar"], n["pay"]) for n in r["neden"]],
                             [("Anlaşma fiyatı farkı", D(7700), D("30.68")), ("Provizyon eksik", D(7500), D("29.88")), ("Eksik belge", D(6000), D("23.90")), ("Mükerrer fatura", D(1800), D("7.17")),
                              ("Poliçe kapsamı dışı", D(1200), D("4.78")), ("Diğer", D(900), D("3.59"))])
            self.assertEqual(([u["tur"] for u in r["uyarilar"]], load_workbook(cikti).sheetnames),
                             (["İtiraz süresi tanımlı değil", "Sınıflandırılamayan kesinti nedeni"], ["Kesinti Takibi", "Şirket Özeti", "Neden Özeti", "Uyarılar"]))

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, k, sr = Path(tmp) / "k.csv", Path(tmp) / "n.csv", Path(tmp) / "s.csv"
            y.write_text("Kurum;Fatura No;Tarih;Tutar;Neden;İtiraz Tarihi;İtiraz Sonucu;Geri Alınan\nA;1;01.10.2026;1.000,00;fiyat;05.10.2026;Kabul;\nA;2;01.10.2026;500;fiyat;yarın;inceleniyor;\n"
                         "A;3;01.10.2026;500;fiyat;;Kısmi kabul;\nA;4;dün;5;x;;;\nA;5;01.10.2026;100;x;02.10.2026;Ret;150\n", encoding="utf-8")
            k.write_text("Kategori;Anahtar Kelimeler;İtiraz Edilebilir\nFiyat;fiyat | TARİFE;\nKapsam;kapsam;Hayır\nBoş;;Evet\n", encoding="utf-8")
            sr.write_text("Şirket;Süre\nA;30\nB;uzun\n", encoding="utf-8")
            self.assertEqual(main.kural_oku(k), [{"kategori": "Fiyat", "kelimeler": ["fiyat", "tarife"], "edilebilir": "Evet"}, {"kategori": "Kapsam", "kelimeler": ["kapsam"], "edilebilir": "Hayır"}])
            l, sure, uy = main.oku(y, sr)
            self.assertEqual(([(x["fatura"], x["itiraz"], x["sonuc"], x["geri"]) for x in l], sure),
                             ([("1", date(2026, 10, 5), "Kabul edildi", D("1000.00")), ("2", None, None, D(0)), ("3", None, "Kısmi kabul", D(0)), ("5", date(2026, 10, 2), "Reddedildi", D(150))], {"a": 30}))
            self.assertEqual(sorted(u["tur"] for u in uy), ["Geri alınan tutar kesintiden büyük", "Kesinti tarihi veya tutarı okunamadı", "Kısmi kabulde geri alınan tutar yok", "Sonuç var, itiraz tarihi yok",
                                                            "İtiraz sonucu tanınmadı", "İtiraz süresi okunamadı", "İtiraz tarihi okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--kesintiler", str(y), "--kurallar", str(k), "--cikti", c, "--tarih", "12.10.2026"]
            self.assertEqual(main.main(arg + ["--sureler", str(sr), "--itiraz-gun", "20", "--yaklasan-gun", "7"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(arg + ["--itiraz-gun", "0"]), 2)
            self.assertEqual(main.main(arg[:-1] + ["dün"]), 2)
            self.assertEqual(main.main(arg + ["--sureler", str(Path(tmp) / "yok.csv")]), 1)
            k.write_text("Kategori;Anahtar Kelimeler\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
