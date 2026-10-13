import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
G = date(2026, 10, 13)


def rn(**kw):
    x = {"tarih": G, "saat": "09:00", "hasta": "A", "kimlik": "12345678950", "brans": "Dahiliye", "hekim": "", "odeme_ham": "SGK", "tip": "SGK", "sirket": "", "police": "", "bitis": None, "bitis_ham": "",
         "provizyon": "", "pdurum": "", "mustahak": "Var", "telefon": "05", "satir": 2}
    return x | kw


def tek(**kw):
    x = main.kontrol([rn(**kw)])["randevular"][0]
    return x["durum"], x["eksikler"]


class KontrolTesti(unittest.TestCase):
    def test_kimlik(self):
        for k in ("12345678950", "23456789060", "34567890170", "45678901280", "56789012390", "67890123400", "78901234510", "10000000146"):
            self.assertTrue(main.kimlik_gecerli(k), k)
        for k in ("12345678901", "02345678950", "1234567895", "123456789500", "1234567895a", "12345678951", "12345678940"):
            self.assertFalse(main.kimlik_gecerli(k), k)
        self.assertEqual((main.maskele("12345678950"), main.maskele("123"), main.kimlik_metni(12345678950.0), main.kimlik_metni(" 123 456 ")), ("123******50", "***", "12345678950", "123456"))

    def test_kurallar(self):
        self.assertEqual(tek(), ("Hazır", []))
        self.assertEqual(tek(kimlik=""), ("Kritik", ["Kimlik no eksik"]))
        self.assertEqual(tek(kimlik="12345678901", telefon=""), ("Kritik", ["Kimlik no hatalı", "Telefon eksik"]))
        self.assertEqual(tek(tip=None, odeme_ham="?"), ("Eksik var", ["Ödeme tipi belirsiz"]))
        self.assertEqual(tek(mustahak="Yok"), ("Kritik", ["SGK müstahaklığı yok"]))
        self.assertEqual(tek(mustahak=""), ("Eksik var", ["Müstahaklık sorgulanmamış"]))
        self.assertEqual(tek(telefon=""), ("Hazır", ["Telefon eksik"]))
        oss = {"tip": "ÖSS", "mustahak": "", "sirket": "S", "police": "P", "bitis": date(2026, 12, 31), "provizyon": "PRV", "pdurum": "Onaylandı"}
        self.assertEqual(tek(**oss), ("Hazır", []))
        self.assertEqual(tek(**oss | {"bitis": G}), ("Hazır", []))                                  # bitiş günü randevu günü: geçerli
        self.assertEqual(tek(**oss | {"bitis": date(2026, 10, 12)}), ("Kritik", ["Poliçe süresi dolmuş"]))
        self.assertEqual(tek(**oss | {"sirket": "", "police": "", "bitis": None, "provizyon": "", "pdurum": ""}),
                         ("Eksik var", ["Sigorta şirketi eksik", "Poliçe no eksik", "Poliçe bitiş tarihi eksik", "Provizyon alınmamış"]))
        self.assertEqual(tek(**oss | {"pdurum": "Reddedildi"}), ("Kritik", ["Provizyon reddedilmiş"]))
        self.assertEqual(tek(**oss | {"pdurum": "RET", "provizyon": ""}), ("Kritik", ["Provizyon reddedilmiş"]))
        self.assertEqual(tek(**oss | {"pdurum": "Bekliyor", "provizyon": ""}), ("Eksik var", ["Provizyon beklemede"]))
        self.assertEqual(tek(**oss | {"tip": "TSS"}), ("Eksik var", ["Müstahaklık sorgulanmamış"]))
        self.assertEqual(tek(**oss | {"tip": "TSS", "mustahak": "Evet"}), ("Hazır", []))
        self.assertEqual(tek(tip="Ücretli", mustahak=""), ("Hazır", []))

    def test_mukerrer_filtre_ozet(self):
        l = [rn(satir=2), rn(satir=3, saat="15:00"), rn(satir=4, brans="Göz"), rn(satir=5, tarih=date(2026, 10, 14), mustahak="Yok"), rn(satir=6, kimlik="", hasta="B", tip="Ücretli"),
             rn(satir=7, kimlik="", hasta="B", tip="Ücretli", saat="16:00")]
        r = main.kontrol(l)
        self.assertEqual([(x["satir"], x["durum"], x["eksikler"]) for x in r["randevular"]],
                         [(2, "Hazır", []), (4, "Hazır", []), (6, "Kritik", ["Kimlik no eksik"]), (3, "Eksik var", ["Mükerrer randevu"]), (7, "Kritik", ["Kimlik no eksik", "Mükerrer randevu"]), (5, "Kritik", ["SGK müstahaklığı yok"])])
        self.assertEqual((r["say"], r["ozet"]), ({"Hazır": 2, "Eksik var": 1, "Kritik": 3}, {"SGK": {"toplam": 4, "Hazır": 2, "Eksik var": 1, "Kritik": 1}, "Ücretli": {"toplam": 2, "Hazır": 0, "Eksik var": 0, "Kritik": 2}}))
        self.assertIn("satır 2", [u for u in r["uyarilar"] if u["tur"] == "Mükerrer randevu"][0]["aciklama"])
        self.assertEqual([x["satir"] for x in main.kontrol(l, date(2026, 10, 14), date(2026, 10, 14))["randevular"]], [5])
        for bas in (date(2026, 11, 1),):
            with self.assertRaises(ValueError):
                main.kontrol(l, bas)
        with self.assertRaises(ValueError):
            main.kontrol([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "p.xlsx"
            r = main.calistir(ORNEK / "randevular.csv", cikti)
            self.assertEqual([x["durum"] for x in r["randevular"]], ["Hazır", "Hazır", "Kritik", "Eksik var", "Kritik", "Kritik", "Eksik var", "Hazır", "Eksik var"])
            self.assertEqual([x["eksikler"] for x in r["randevular"]][2:7], [["Poliçe süresi dolmuş", "Provizyon alınmamış"], ["Müstahaklık sorgulanmamış", "Provizyon beklemede"], ["SGK müstahaklığı yok"],
                                                                             ["Kimlik no hatalı"], ["Ödeme tipi belirsiz", "Telefon eksik"]])
            self.assertEqual((r["say"], r["ozet"]["Belirsiz"]["toplam"], len(r["uyarilar"])), ({"Hazır": 3, "Eksik var": 3, "Kritik": 3}, 1, 9))
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Kontrol Listesi", "Özet", "Uyarılar"])
            self.assertEqual([c.value for c in wb["Kontrol Listesi"][2]][:5], ["Hazır", wb["Kontrol Listesi"]["B2"].value, "09:00", "Örnek Hasta Bir", "123******50"])
            self.assertNotIn("12345678950", str([[c.value for c in row] for ws in wb for row in ws.iter_rows()]))

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "r.csv"
            y.write_text("Tarih;Hasta Adı;TCKN;Güvence;Bölüm\n13.10.2026;A;12345678950;Özel Sağlık Sigortası;Göz\nyarın;B;1;SGK;Göz\n;;;;\n2026-10-14;C;23456789060;Nakit;Göz\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["hasta"], x["tip"], x["tarih"]) for x in l], [u["tur"] for u in uy]), ([("A", "ÖSS", G), ("C", "Ücretli", date(2026, 10, 14))], ["Randevu okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--randevular", str(y), "--cikti", c]
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(arg + ["--baslangic", "14.10.2026", "--bitis", "14.10.2026"]), 0)
            self.assertEqual(main.main(arg + ["--baslangic", "01.11.2026"]), 1)
            self.assertEqual(main.main(arg + ["--baslangic", "yarın"]), 2)
            self.assertEqual(main.main(arg + ["--baslangic", "14.10.2026", "--bitis", "13.10.2026"]), 2)
            self.assertEqual(main.main(["--randevular", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Hasta;Saat\nA;1\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
