"""
Provizyon ve Sigorta Bilgisi Ön Kontrolü — Workers / Workless kod bloğu
Hasta Hizmetleri › Hasta Kayıt Kabul Görevlisi

Randevu listesindeki hastaların kimlik, ödeme tipi, sigorta ve provizyon bilgilerini hasta gelmeden önce kontrol edip eksikleri listeler.
  Kimlik no        = 11 hane, ilk hane 0 olamaz; 10. hane = ((1,3,5,7,9. haneler toplamı × 7) − (2,4,6,8. haneler toplamı)) mod 10; 11. hane = ilk 10 hane toplamı mod 10
  SGK, TSS         = müstahaklık "Yok" → kritik; boş → sorgulanmamış
  ÖSS, TSS         = sigorta şirketi, poliçe no, poliçe bitiş tarihi eksik mi; poliçe bitişi randevu tarihinden önce mi (kritik)
                     provizyon reddedilmiş (kritik) · beklemede · alınmamış
  Diğer            = ödeme tipi belirsiz · telefon eksik · aynı gün aynı branşta mükerrer randevu
  Durum            = Kritik (yüksek önemli eksik var) · Eksik var · Hazır
Kod Medula'ya veya sigorta şirketine bağlanmaz; yalnız listenizdeki bilgiyi denetler. Müstahaklık ve provizyon sonucunu sorgulayıp listeye siz yazarsınız.
Kimlik numaraları çıktıda maskelenir. Liste özel nitelikli kişisel veri içerir; KVKK kapsamında yetkisiz kişilerle paylaşmayın.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"tarih": ("randevu tarihi", "tarih"), "saat": ("saat", "randevu saati"), "hasta": ("hasta", "hasta adi", "ad soyad", "hasta adi soyadi"),
         "kimlik": ("tc kimlik no", "tckn", "kimlik no", "tc no", "tc"), "brans": ("brans", "bolum", "poliklinik"), "hekim": ("hekim", "doktor"),
         "odeme": ("odeme tipi", "kurum tipi", "odeme turu", "guvence"), "sirket": ("sigorta sirketi", "kurum", "sirket"), "police": ("police no", "police", "sigortali no"),
         "bitis": ("police bitis", "police bitis tarihi", "bitis tarihi"), "provizyon": ("provizyon no", "provizyon", "onay no"), "pdurum": ("provizyon durumu", "onay durumu"),
         "mustahak": ("mustahaklik", "sgk mustahaklik", "mustehaklik"), "telefon": ("telefon", "cep telefonu", "gsm")}
TIP = {"sgk": "SGK", "oss": "ÖSS", "ozel saglik sigortasi": "ÖSS", "ozel sigorta": "ÖSS", "tss": "TSS", "tamamlayici": "TSS", "tamamlayici saglik sigortasi": "TSS",
       "ucretli": "Ücretli", "nakit": "Ücretli", "ozel": "Ücretli", "cepten": "Ücretli"}
VAR, YOK = {"var", "evet", "aktif", "e", "uygun"}, {"yok", "hayir", "pasif", "h"}
DURUM_RENK = {"Kritik": KIRMIZI, "Eksik var": SARI, "Hazır": YESIL}


def kimlik_metni(x) -> str:
    if isinstance(x, float) and x == int(x):
        x = int(x)
    return re.sub(r"\s", "", metin(x))


def kimlik_gecerli(k: str) -> bool:
    if not re.fullmatch(r"[1-9]\d{10}", k):
        return False
    d = [int(c) for c in k]
    return (sum(d[0:9:2]) * 7 - sum(d[1:8:2])) % 10 == d[9] and sum(d[:10]) % 10 == d[10]


def maskele(k: str) -> str:
    return k[:3] + "*" * (len(k) - 5) + k[-2:] if len(k) > 5 else "*" * len(k)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("tarih", "hasta", "odeme")):
        h = metin(r["hasta"])
        if not h:
            continue
        t = tarih(r["tarih"])
        if t is None:
            uy.append(uyari("Yüksek", "Randevu okunamadı", f"{h} (satır {r['_satir']})", f"Randevu tarihi '{metin(r['tarih'])}' okunamadı. Satır alınmadı."))
            continue
        l.append({"tarih": t, "saat": metin(r["saat"])[:5], "hasta": h, "kimlik": kimlik_metni(r["kimlik"]), "brans": metin(r["brans"]), "hekim": metin(r["hekim"]), "odeme_ham": metin(r["odeme"]),
                  "tip": TIP.get(katla(r["odeme"])), "sirket": metin(r["sirket"]), "police": metin(r["police"]), "bitis": tarih(r["bitis"]), "bitis_ham": metin(r["bitis"]),
                  "provizyon": metin(r["provizyon"]), "pdurum": metin(r["pdurum"]), "mustahak": metin(r["mustahak"]), "telefon": metin(r["telefon"]), "satir": r["_satir"]})
    return l, uy


def kontrol(randevular: list[dict], bas: date | None = None, bit: date | None = None) -> dict:
    secili = [x for x in randevular if (bas is None or x["tarih"] >= bas) and (bit is None or x["tarih"] <= bit)]
    if not secili:
        raise ValueError("Kontrol edilecek randevu yok (dosya boş veya seçilen tarih aralığında randevu bulunmuyor).")
    gorulen, liste, uy = {}, [], []
    for x in sorted(secili, key=lambda x: (x["tarih"], x["saat"], x["satir"])):
        e, k, tip = [], x["kimlik"], x["tip"]
        if not k:
            e.append(("Yüksek", "Kimlik no eksik", "Kimlik numarası girilmemiş; müstahaklık ve provizyon sorgulanamaz."))
        elif not kimlik_gecerli(k):
            e.append(("Yüksek", "Kimlik no hatalı", f"'{maskele(k)}' 11 haneli değil veya doğrulama haneleri tutmuyor."))
        if tip is None:
            e.append(("Orta", "Ödeme tipi belirsiz", f"Ödeme tipi '{x['odeme_ham']}' tanınmadı (SGK, ÖSS, TSS veya Ücretli olmalı)."))
        if tip in ("SGK", "TSS"):
            m = katla(x["mustahak"])
            if m in YOK:
                e.append(("Yüksek", "SGK müstahaklığı yok", "Müstahaklık sorgusu olumsuz. Hastayı arayıp ücretli muayene veya güvence güncellemesi konusunda bilgilendirin."))
            elif m not in VAR:
                e.append(("Orta", "Müstahaklık sorgulanmamış", "SGK müstahaklık sonucu listede yok."))
        if tip in ("ÖSS", "TSS"):
            if not x["sirket"]:
                e.append(("Orta", "Sigorta şirketi eksik", "Sigorta şirketi yazılmamış."))
            if not x["police"]:
                e.append(("Orta", "Poliçe no eksik", "Poliçe veya sigortalı numarası yazılmamış."))
            if x["bitis"] is None:
                e.append(("Orta", "Poliçe bitiş tarihi eksik", f"Poliçe bitiş tarihi '{x['bitis_ham']}' okunamadı veya boş."))
            elif x["bitis"] < x["tarih"]:
                e.append(("Yüksek", "Poliçe süresi dolmuş", f"Poliçe {x['bitis']:%d.%m.%Y} tarihinde bitiyor, randevu {x['tarih']:%d.%m.%Y}. Yenilenen poliçeyi isteyin."))
            pd = katla(x["pdurum"])
            if re.search(r"\bre[dt]", pd):
                e.append(("Yüksek", "Provizyon reddedilmiş", f"Provizyon durumu: '{x['pdurum']}'. Ret gerekçesini öğrenip hastayı bilgilendirin."))
            elif "bekl" in pd:
                e.append(("Orta", "Provizyon beklemede", f"Provizyon durumu: '{x['pdurum']}'. Randevudan önce sonucu takip edin."))
            elif not x["provizyon"]:
                e.append(("Orta", "Provizyon alınmamış", "Provizyon numarası yok. Ön onay gerektiren işlemlerde randevudan önce başvurun."))
        if not x["telefon"]:
            e.append(("Bilgi", "Telefon eksik", "Hastaya ulaşılacak telefon yok."))
        a = (k or katla(x["hasta"]), x["tarih"], katla(x["brans"]))
        if a in gorulen:
            e.append(("Orta", "Mükerrer randevu", f"Aynı gün aynı branşta başka randevu var (satır {gorulen[a]})."))
        else:
            gorulen[a] = x["satir"]
        onemler = {o for o, _, _ in e}
        durum = "Kritik" if "Yüksek" in onemler else "Eksik var" if "Orta" in onemler else "Hazır"
        liste.append(x | {"durum": durum, "eksikler": [t for _, t, _ in e], "maske": maskele(k) if k else ""})
        uy += [uyari(o, t, f"{x['hasta']} · {x['tarih']:%d.%m.%Y} {x['saat']}".strip(), ac) for o, t, ac in e]
    ozet = {}
    for x in liste:
        o = ozet.setdefault(x["tip"] or "Belirsiz", {"toplam": 0, "Hazır": 0, "Eksik var": 0, "Kritik": 0})
        o["toplam"] += 1
        o[x["durum"]] += 1
    say = {d: sum(1 for x in liste if x["durum"] == d) for d in ("Hazır", "Eksik var", "Kritik")}
    return {"randevular": liste, "ozet": ozet, "say": say, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kontrol Listesi"
    _baslik(ws, ["Durum", "Tarih", "Saat", "Hasta", "Kimlik No (maskeli)", "Branş", "Hekim", "Ödeme Tipi", "Sigorta Şirketi", "Eksikler", "Telefon", "Not (siz doldurun)"],
            (11, 12, 7, 26, 16, 18, 22, 11, 22, 52, 16, 30))
    for x in r["randevular"]:
        _satir(ws, [x["durum"], x["tarih"], x["saat"], x["hasta"], x["maske"], x["brans"], x["hekim"], x["tip"] or x["odeme_ham"], x["sirket"], ", ".join(x["eksikler"]), x["telefon"], ""],
               tarih_s=(2,), dolgu={1: DURUM_RENK[x["durum"]]}, sar=(10,))
    ws = wb.create_sheet("Özet")
    _baslik(ws, ["Ödeme Tipi", "Randevu", "Hazır", "Eksik Var", "Kritik"], (16, 10, 10, 10, 10))
    for tip, o in sorted(r["ozet"].items()):
        _satir(ws, [tip, o["toplam"], o["Hazır"], o["Eksik var"], o["Kritik"]], dolgu={5: KIRMIZI if o["Kritik"] else None})
    _satir(ws, ["Toplam", len(r["randevular"]), r["say"]["Hazır"], r["say"]["Eksik var"], r["say"]["Kritik"]])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.append([])
    ws.append(["Kod yalnız listedeki bilgiyi denetler; Medula'ya veya sigorta şirketine bağlanmaz. Liste kişisel sağlık verisi içerir, yetkisiz kişilerle paylaşmayın."])
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bas: date | None = None, bit: date | None = None) -> dict:
    randevular, uy = oku(yol)
    r = kontrol(randevular, bas, bit)
    uy += r["uyarilar"]
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Randevu listesindeki kimlik, sigorta ve provizyon bilgilerini ön kontrolden geçirir.")
    p.add_argument("--randevular", type=Path, default=ORNEK / "randevular.csv",
                   help="Randevu Tarihi, Saat, Hasta, TC Kimlik No, Branş, Hekim, Ödeme Tipi, Sigorta Şirketi, Poliçe No, Poliçe Bitiş, Provizyon No, Provizyon Durumu, Müstahaklık, Telefon")
    p.add_argument("--baslangic", help="Yalnız bu tarihten itibaren olan randevular (GG.AA.YYYY)")
    p.add_argument("--bitis", help="Yalnız bu tarihe kadar olan randevular (GG.AA.YYYY)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "provizyon_on_kontrol.xlsx")
    a = p.parse_args(argv)
    bas, bit = (tarih(a.baslangic) if a.baslangic else None), (tarih(a.bitis) if a.bitis else None)
    if (a.baslangic and bas is None) or (a.bitis and bit is None) or (bas and bit and bit < bas):
        print("[X] --baslangic ve --bitis GG.AA.YYYY biçiminde olmalı; bitiş başlangıçtan önce olamaz")
        return 2
    if not a.randevular.exists():
        print(f"[X] Dosya bulunamadı: {a.randevular}")
        return 1
    try:
        r = calistir(a.randevular, a.cikti, bas, bit)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    s = r["say"]
    print(f"[OK] {len(r['randevular'])} randevu · {s['Hazır']} hazır · {s['Eksik var']} eksik var · {s['Kritik']} kritik")
    for x in r["randevular"]:
        if x["durum"] == "Kritik":
            print(f"[!] {x['tarih']:%d.%m.%Y} {x['saat']} {x['hasta']}: {', '.join(x['eksikler'])}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
