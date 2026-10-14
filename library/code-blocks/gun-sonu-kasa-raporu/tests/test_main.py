import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def k(kasa, nakit, kart=0, diger=0, n_iade=0, k_iade=0, iptal=0, cikis=0, sayilan=None, pos=None, acilis=500, fis=None, iptal_adet=None):
    c = lambda x: None if x is None else D(str(x))   # noqa: E731
    return {"kasa": kasa, "kasiyer": "", "acilis": D(acilis), "nakit": D(nakit), "kart": D(kart), "diger": D(diger), "n_iade": D(n_iade), "k_iade": D(k_iade), "iptal": D(iptal), "cikis": D(cikis), "sayilan": c(sayilan), "pos": c(pos),
            "iptal_adet": c(iptal_adet), "fis": c(fis)}


def veri():
    return [k("K1", 10000, 20000, 1000, 300, 500, 400, 2000, 8195, 19500, fis=200, iptal_adet=2), k("K2", 5000, 8000, sayilan=5380, pos=7900, fis=100), k("K3", 3000), k("K4", 3000, sayilan=4100)]


class RaporTesti(unittest.TestCase):
    def test_mutabakat(self):
        r = main.raporla(veri())
        self.assertEqual({x["kasa"]: (x["durum"], x["beklenen"], x["n_fark"], x["net_kart"], x["k_fark"], x["brut"], x["iade"], x["net"], x["iade_oran"], x["iptal_oran"], x["teslim"]) for x in r["kasalar"]},
                         {"K1": ("Tamam", D(8200), D(-5), D(19500), D(0), D(31000), D(800), D(30200), D("2.58"), D("1.29"), D(7695)),
                          "K2": ("Nakit eksiği; Kart farkı", D(5500), D(-120), D(8000), D(-100), D(13000), D(0), D(13000), D("0.00"), D("0.00"), D(4880)),
                          "K3": ("Sayım yok", D(3500), None, D(0), None, D(3000), D(0), D(3000), D("0.00"), D("0.00"), None), "K4": ("Nakit fazlası", D(3500), D(600), D(0), None, D(3000), D(0), D(3000), D("0.00"), D("0.00"), D(3600))})
        t = r["toplam"]
        self.assertEqual((t["kasa"], t["nakit"], t["kart"], t["diger"], t["brut"], t["iade"], t["net"], t["beklenen"], t["sayilan"], t["n_fark"], t["k_fark"], t["pos"], t["teslim"], t["fis"], t["iptal_adet"], t["farkli"], t["sepet"], t["iptal_oran"]),
                         (4, D(21000), D(28000), D(1000), D(50000), D(800), D(49200), D(20700), D(17675), D(475), D(-100), D(27400), D(16175), 300, 2, 3, D("164.00"), D("0.80")))
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Orta", "Kart tutarı POS gün sonuyla uyuşmuyor", "K2"), ("Orta", "Kasa eksiği", "K2"), ("Orta", "Kasa sayımı girilmemiş", "K3"), ("Yüksek", "Kasa fazlası", "K4")])

    def test_tolerans(self):
        d = lambda sayilan, **kw: main.raporla([k("K", 1000, sayilan=sayilan)], **kw)   # noqa: E731
        self.assertEqual([d(x)["kasalar"][0]["durum"] for x in (1495, 1494.99, 1505, 1505.01)], ["Tamam", "Nakit eksiği", "Tamam", "Nakit fazlası"])
        self.assertEqual(([u["onem"] for u in d(1000.01)["uyarilar"]], [u["onem"] for u in d(1000)["uyarilar"]], d(1400, tolerans=D(100))["uyarilar"], [u["onem"] for u in d(1400, tolerans=D(0), buyuk=D(100))["uyarilar"]]),
                         (["Orta"], ["Yüksek"], [], ["Yüksek"]))
        r = main.raporla([k("K", 100, cikis=900, sayilan=0, pos=0, kart=50)], buyuk=D(50))
        self.assertEqual(sorted((u["onem"], u["tur"]) for u in r["uyarilar"]), [("Orta", "Beklenen nakit eksi"), ("Yüksek", "Kart tutarı POS gün sonuyla uyuşmuyor"), ("Yüksek", "Kasa fazlası")])
        with self.assertRaises(ValueError):
            main.raporla([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "kasa_gun_sonu.csv", cikti)
            self.assertEqual([(x["kasa"], x["durum"], x["beklenen"], x["n_fark"], x["net_kart"], x["k_fark"], x["net"], x["iptal_oran"]) for x in r["kasalar"]],
                             [("Kasa 1", "Tamam", D("9130.50"), D("-2.00"), D(41410), D(0), D("60790.50"), D("1.03")), ("Kasa 2", "Nakit eksiği", D(8980), D(-150), D(35310), D(0), D(49090), D("3.67")),
                              ("Kasa 3", "Kart farkı", D(10460), D(0), D(27150), D(-250), D(36610), D("0.57"))])
            t = r["toplam"]
            self.assertEqual((t["nakit"], t["kart"], t["diger"], t["brut"], t["iade"], t["net"], t["n_fark"], t["k_fark"], t["fis"], t["sepet"], t["farkli"], t["teslim"]),
                             (D("41070.50"), D(105210), D(2050), D("148330.50"), D(1840), D("146490.50"), D("-152.00"), D(-250), 777, D("188.53"), 2, D("25418.50")))
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]], [("Orta", "Kart tutarı POS gün sonuyla uyuşmuyor", "Kasa 3 · C. Kaya"), ("Orta", "Kasa eksiği", "Kasa 2 · B. Demir")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Kasa Mutabakatı", "Satış Özeti", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "k.csv"
            y.write_text("Kasa No;Personel;Devir Nakit;Nakit;Kart;İade;İptal;Bankaya Yatan;Sayım;POS Toplamı;Fiş Sayısı\nK1;Ali;500;1.250,50;2.000;50;10;300;1.400,50;2.000;40\nK2;;;100;;;;;;;\nk1;;;5;;;;;;;\nK3;;;x;;;;;;;\n"
                         "K4;;;10;-5;;;;;;\nK5;;;10;;;;;y;;\n;;;10;;;;;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["kasa"], x["kasiyer"], x["acilis"], x["nakit"], x["kart"], x["n_iade"], x["k_iade"], x["iptal"], x["cikis"], x["sayilan"], x["pos"], x["fis"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("K1", "Ali", D(500), D("1250.50"), D(2000), D(50), D(0), D(10), D(300), D("1400.50"), D(2000), D(40)), ("K2", "", D(0), D(100), D(0), D(0), D(0), D(0), D(0), None, None, None)],
                              [("Kasa satırı okunamadı", x) for x in ("k1", "K3", "K4", "K5")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--kasalar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tolerans", "10", "--buyuk", "250"]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tolerans", "-1"], ["--buyuk", "5"], ["--buyuk", "x"])], [2] * 3)
            self.assertEqual(main.main(["--kasalar", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Kasa;Kasiyer\nA;B\n", "Kasa;Nakit Satış\n;5\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
