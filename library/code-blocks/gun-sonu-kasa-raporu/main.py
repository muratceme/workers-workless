"""
Gün Sonu Kasa Raporu — Gerisi Bende kod bloğu
Mağaza Operasyonları › Kasa Sorumlusu

Kasa bazında Z raporu ve POS toplamlarından nakit, kredi kartı, iade ve iptal toplamlarını çıkarır; kasa sayımıyla farkı raporlar.
  Beklenen nakit      = açılış nakdi + nakit satış − nakit iade − kasadan çıkış
  Nakit farkı         = sayılan nakit − beklenen nakit           (eksi = kasa eksiği, artı = kasa fazlası)
  Net kart satışı     = kredi kartı satış − kart iade
  Kart farkı          = POS gün sonu toplamı − net kart satışı   (POS gün sonu yazılmışsa)
  Brüt satış          = nakit + kredi kartı + diğer ödemeler;   net satış = brüt satış − iadeler
  İade % / iptal %    = iade ya da iptal tutarı ÷ brüt satış × 100
  Teslim edilecek nakit = sayılan nakit − açılış nakdi           (açılış nakdi kadar tutar kasada bırakılırsa)
  Durum               = fark --tolerans (5) içindeyse Tamam; değilse Nakit eksiği / Nakit fazlası / Kart farkı
Tutarları Z raporundan ve POS gün sonu fişinden siz girersiniz; kod yazar kasaya ya da bankaya bağlanmaz.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"kasa": ("kasa", "kasa no", "kasa adi"), "kasiyer": ("kasiyer", "kasa sorumlusu", "personel"), "acilis": ("acilis nakdi", "devir nakit", "acilis", "kasa acilis"), "nakit": ("nakit satis", "nakit"),
         "kart": ("kredi karti satis", "kart satis", "kredi karti", "kart"), "diger": ("diger odeme", "diger odemeler", "yemek karti", "diger"), "n_iade": ("nakit iade", "iade nakit", "iade"),
         "k_iade": ("kart iade", "kredi karti iade", "iade kart"), "iptal": ("iptal tutari", "iptal"), "iptal_adet": ("iptal adedi", "iptal sayisi"), "cikis": ("kasadan cikis", "kasa cikis", "bankaya yatan", "tediye"),
         "sayilan": ("sayilan nakit", "kasa sayimi", "sayim"), "pos": ("pos gun sonu", "pos toplami", "pos slip toplami", "banka gun sonu"), "fis": ("fis adedi", "fis sayisi", "islem adedi")}
TUTAR = ("acilis", "nakit", "kart", "diger", "n_iade", "k_iade", "iptal", "cikis")
SIFIR = Decimal(0)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("kasa", "nakit")):
        kasa = metin(r["kasa"])
        if not kasa:
            continue
        d = {a: para(r[a]) if metin(r[a]) else SIFIR for a in TUTAR}
        ek = {a: para(r[a]) if metin(r[a]) else None for a in ("sayilan", "pos", "iptal_adet", "fis")}
        if any(v is None or v < 0 for v in d.values()) or any(metin(r[a]) and (ek[a] is None or ek[a] < 0) for a in ek) or katla(kasa) in gorulen:
            uy.append(uyari("Yüksek", "Kasa satırı okunamadı", kasa, f"Satır {r['_satir']}: " + ("aynı kasa daha önce geçti" if katla(kasa) in gorulen else "tutarlardan biri okunamadı ya da eksi") + ". Satır alınmadı."))
            continue
        gorulen.add(katla(kasa))
        l.append({"kasa": kasa, "kasiyer": metin(r["kasiyer"]), **d, **ek})
    return l, uy


def _oran(pay: Decimal, payda: Decimal) -> Decimal | None:
    return yuvarla(pay * 100 / payda) if payda else None


def raporla(kasalar: list[dict], tolerans: Decimal = Decimal(5), buyuk: Decimal = Decimal(500)) -> dict:
    if not kasalar:
        raise ValueError("Raporlanacak kasa yok.")
    uy = []
    for k in kasalar:
        kim = f"{k['kasa']} · {k['kasiyer']}" if k["kasiyer"] else k["kasa"]
        k["beklenen"] = k["acilis"] + k["nakit"] - k["n_iade"] - k["cikis"]
        k["n_fark"] = k["sayilan"] - k["beklenen"] if k["sayilan"] is not None else None
        k["net_kart"] = k["kart"] - k["k_iade"]
        k["k_fark"] = k["pos"] - k["net_kart"] if k["pos"] is not None else None
        k["brut"], k["iade"] = k["nakit"] + k["kart"] + k["diger"], k["n_iade"] + k["k_iade"]
        k["net"] = k["brut"] - k["iade"]
        k["iade_oran"], k["iptal_oran"] = _oran(k["iade"], k["brut"]), _oran(k["iptal"], k["brut"])
        k["teslim"] = k["sayilan"] - k["acilis"] if k["sayilan"] is not None else None
        durum = []
        if k["sayilan"] is None:
            durum.append("Sayım yok")
            uy.append(uyari("Orta", "Kasa sayımı girilmemiş", kim, f"Beklenen nakit {k['beklenen']:,.2f}; sayılan nakit yazılmadığı için fark hesaplanamadı."))
        elif abs(k["n_fark"]) > tolerans:
            durum.append("Nakit eksiği" if k["n_fark"] < 0 else "Nakit fazlası")
            uy.append(uyari("Yüksek" if abs(k["n_fark"]) >= buyuk else "Orta", "Kasa eksiği" if k["n_fark"] < 0 else "Kasa fazlası", kim, f"Beklenen nakit {k['beklenen']:,.2f}, sayılan {k['sayilan']:,.2f}; fark {k['n_fark']:+,.2f}. "
                            "Sayımı tekrarlayın; iade, iptal ve kasadan çıkış fişlerini kontrol edin."))
        if k["k_fark"] is not None and abs(k["k_fark"]) > tolerans:
            durum.append("Kart farkı")
            uy.append(uyari("Yüksek" if abs(k["k_fark"]) >= buyuk else "Orta", "Kart tutarı POS gün sonuyla uyuşmuyor", kim, f"Yazar kasada net kart satışı {k['net_kart']:,.2f}, POS gün sonu {k['pos']:,.2f}; fark {k['k_fark']:+,.2f}. "
                            "Ödeme türü yanlış seçilmiş ya da slip iptali işlenmemiş olabilir."))
        if k["beklenen"] < 0:
            uy.append(uyari("Orta", "Beklenen nakit eksi", kim, f"Açılış + nakit satış − iade − çıkış = {k['beklenen']:,.2f}. Girilen tutarları kontrol edin."))
        k["durum"] = "; ".join(durum) or "Tamam"
    t = {a: sum((k[a] for k in kasalar), SIFIR) for a in TUTAR + ("beklenen", "net_kart", "brut", "iade", "net")}
    sayilan = [k for k in kasalar if k["sayilan"] is not None]
    t.update(kasa=len(kasalar), sayilan=sum((k["sayilan"] for k in sayilan), SIFIR), n_fark=sum((k["n_fark"] for k in sayilan), SIFIR), k_fark=sum((k["k_fark"] for k in kasalar if k["k_fark"] is not None), SIFIR),
             pos=sum((k["pos"] for k in kasalar if k["pos"] is not None), SIFIR), teslim=sum((k["teslim"] for k in sayilan), SIFIR), fis=sum((int(k["fis"]) for k in kasalar if k["fis"] is not None), 0),
             iptal_adet=sum((int(k["iptal_adet"]) for k in kasalar if k["iptal_adet"] is not None), 0), iade_oran=_oran(t["iade"], t["brut"]), iptal_oran=_oran(t["iptal"], t["brut"]), farkli=sum(1 for k in kasalar if k["durum"] != "Tamam"))
    t["sepet"] = yuvarla(t["net"] / t["fis"]) if t["fis"] else None
    return {"kasalar": kasalar, "toplam": t, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, tolerans: Decimal) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kasa Mutabakatı"
    _baslik(ws, ["Kasa", "Kasiyer", "Durum", "Açılış Nakdi", "Nakit Satış", "Nakit İade", "Kasadan Çıkış", "Beklenen Nakit", "Sayılan Nakit", "Nakit Farkı", "Net Kart Satışı", "POS Gün Sonu", "Kart Farkı", "Teslim Edilecek Nakit"],
            (12, 20, 26, 13, 14, 12, 13, 14, 14, 12, 14, 14, 12, 14), "B2")
    t = r["toplam"]
    for k in r["kasalar"] + [{"kasa": "TOPLAM", "kasiyer": None, "durum": None, **t}]:
        n = _satir(ws, [k["kasa"], k["kasiyer"], k["durum"], k["acilis"], k["nakit"], k["n_iade"], k["cikis"], k["beklenen"], k["sayilan"], k["n_fark"], k["net_kart"], k["pos"], k["k_fark"], k["teslim"]], para_s=range(4, 15),
                   dolgu={3: None if k["durum"] is None else YESIL if k["durum"] == "Tamam" else KIRMIZI, 10: KIRMIZI if k["durum"] and k["n_fark"] is not None and abs(k["n_fark"]) > tolerans else None,
                          13: KIRMIZI if k["durum"] and k["k_fark"] is not None and abs(k["k_fark"]) > tolerans else None})
    for h in ws[n]:
        h.font = Font(bold=True)
    ws.append([])
    ws.append([f"Fark toleransı {tolerans}. Nakit farkı eksi ise kasa eksiği, artı ise kasa fazlasıdır."])
    ws = wb.create_sheet("Satış Özeti")
    _baslik(ws, ["Kasa", "Kasiyer", "Nakit Satış", "Kredi Kartı Satış", "Diğer Ödeme", "Brüt Satış", "Nakit İade", "Kart İade", "Net Satış", "İade %", "İptal Tutarı", "İptal Adedi", "İptal %", "Fiş Adedi"],
            (12, 20, 14, 14, 13, 14, 12, 12, 14, 9, 12, 9, 9, 9), "B2")
    for k in r["kasalar"] + [{"kasa": "TOPLAM", "kasiyer": None, **t}]:
        n = _satir(ws, [k["kasa"], k["kasiyer"], k["nakit"], k["kart"], k["diger"], k["brut"], k["n_iade"], k["k_iade"], k["net"], k["iade_oran"], k["iptal"], k["iptal_adet"], k["iptal_oran"], k["fis"]], para_s=(3, 4, 5, 6, 7, 8, 9, 10, 11, 13))
    for h in ws[n]:
        h.font = Font(bold=True)
    ws.append([])
    ws.append(["Fiş başına net satış", float(t["sepet"]) if t["sepet"] is not None else None])
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, tolerans: Decimal = Decimal(5), buyuk: Decimal = Decimal(500)) -> dict:
    kasalar, uy = oku(yol)
    r = raporla(kasalar, tolerans, buyuk)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, tolerans)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="POS ve yazar kasa Z raporlarından nakit, kredi kartı, iade ve iptal toplamlarını çıkarır, kasa sayımıyla farkı raporlar.")
    p.add_argument("--kasalar", type=Path, help="Kasa, Kasiyer, Açılış Nakdi, Nakit Satış, Kredi Kartı Satış, Diğer Ödeme, Nakit İade, Kart İade, İptal Tutarı, İptal Adedi, Kasadan Çıkış, Sayılan Nakit, POS Gün Sonu, Fiş Adedi")
    p.add_argument("--tolerans", default="5", help="Bu tutara kadar fark 'Tamam' sayılır (varsayılan 5)")
    p.add_argument("--buyuk", default="500", help="Bu tutar ve üstü fark yüksek önemli sayılır (varsayılan 500)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "gun_sonu_kasa_raporu.xlsx")
    a = p.parse_args(argv)
    yol = a.kasalar or ORNEK / "kasa_gun_sonu.csv"
    tol, buyuk = para(a.tolerans), para(a.buyuk)
    if tol is None or tol < 0 or buyuk is None or buyuk <= tol:
        print("[X] --tolerans eksi olamaz; --buyuk toleranstan büyük olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, tol, buyuk)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {t['kasa']} kasa · net satış {t['net']:,.2f} (nakit {t['nakit']:,.2f}, kart {t['kart']:,.2f}, diğer {t['diger']:,.2f}) · iade {t['iade']:,.2f} · iptal {t['iptal']:,.2f} · nakit farkı {t['n_fark']:+,.2f} · farklı kasa {t['farkli']}")
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
