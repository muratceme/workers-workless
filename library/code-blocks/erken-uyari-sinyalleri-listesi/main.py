"""
Erken Uyarı Sinyalleri Listesi — Workers / Workless kod bloğu
Bankacılık › Krediler İzleme ve Takip › Kredi İzleme Uzmanı

Kredi müşterilerine ait sinyalleri tek listede birleştirir ve firma bazında risk puanı üretir:
  - Gecikme günü (portföy dökümünden)
  - Limit aşımı / limit doluluğu (risk ÷ limit)
  - Karşılıksız çek ve protestolu senet kayıtları (KKB / Risk Merkezi dökümü)
  - Haciz ihbarnameleri (kamu / özel)
  - Ciro düşüşü (son 3 ay ↔ geçen yılın aynı 3 ayı; yoksa önceki 3 ay)
Puanlar ve eşikler ayar dosyasındadır (örnek değerler); kurumunuzun kredi politikasına göre değiştirin.
Araç karar vermez: kredi izleme biriminin inceleme sırasını belirler. İnternete bağlanmaz.

Kullanım:
    python main.py                                         # örnek verilerle dener
    python main.py --portfoy portfoy.xlsx --cek-senet kkb.xlsx --haciz haciz.xlsx --ciro ciro.xlsx --tarih 30.09.2026
    python main.py ... --ayar ayarlar.json
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
SIFIR = Decimal(0)

VARSAYILAN_AYAR = {
    "pencere_ay": 12,
    "gecikme": [[1, 10], [31, 25], [61, 40], [91, 60]],
    "karsiliksiz_cek": {"acik": 15, "odenmis": 5, "ust_sinir": 30},
    "protestolu_senet": {"acik": 10, "odenmis": 3, "ust_sinir": 25},
    "haciz": {"kamu": 15, "ozel": 12, "ust_sinir": 30},
    "limit": {"asim": 15, "doluluk_esik": 95, "doluluk": 5},
    "ciro": [[30, 10], [50, 20]],
    "siniflar": [[70, "Kritik"], [40, "Yakın İzleme"], [20, "İzleme"], [0, "Normal"]],
}


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def katla(s) -> str:
    s = kucuk(s)
    for a, b in zip("ıöüşğç", "iousgc"):
        s = s.replace(a, b)
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def para(x) -> Decimal:
    if x in (None, ""):
        return SIFIR
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("₺", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1 or re.fullmatch(r"-?\d{1,3}\.\d{3}", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return SIFIR


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            continue
    return None


def donem(x) -> tuple[int, int] | None:
    """'2026-09', '09.2026', '2026/09', '01.09.2026', tarih nesnesi → (yıl, ay)."""
    t = tarih(x)
    if t:
        return t.year, t.month
    s = str(x or "").strip()
    m = re.fullmatch(r"(\d{4})[-/.](\d{1,2})", s) or None
    if m:
        return int(m[1]), int(m[2])
    m = re.fullmatch(r"(\d{1,2})[-/.](\d{4})", s)
    if m:
        return int(m[2]), int(m[1])
    return None


def ay_kaydir(d: tuple[int, int], n: int) -> tuple[int, int]:
    i = d[0] * 12 + d[1] - 1 + n
    return i // 12, i % 12 + 1


def tl(x: Decimal) -> str:
    return f"{x:,.0f}".replace(",", ".")


def tablo_oku(yol: Path, anahtarlar=("musteri no", "firma no", "vkn", "vergi no", "musteri")) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    satirlar = [r for r in satirlar if any(c not in (None, "") for c in r)]
    bi = next((i for i, r in enumerate(satirlar[:10]) if any(katla(c) in anahtarlar for c in r)), 0)
    return satirlar[bi:]


def _bul(baslik, *adlar):
    b = [katla(x) for x in baslik]
    return next((b.index(katla(a)) for a in adlar if katla(a) in b), None)


def _al(r, i):
    return r[i] if i is not None and i < len(r) else None


MUSTERI = ("müşteri no", "firma no", "vkn", "vergi no", "müşteri")


def anahtar(x) -> str:
    return re.sub(r"[^0-9a-z]", "", katla(x))


def _sutunlar(yol: Path, tanim: dict, zorunlu: tuple) -> tuple[list[list], dict]:
    s = tablo_oku(yol)
    i = {k: _bul(s[0], *v) for k, v in tanim.items()}
    eksik = [k for k in zorunlu if i[k] is None]
    if eksik:
        raise SystemExit(f"{yol.name}: gerekli sütun bulunamadı ({', '.join(eksik)}). Başlıklar: {s[0]}")
    return s[1:], i


# ----------------------------------------------------------------------------
# Okuma
# ----------------------------------------------------------------------------

def portfoy_oku(yol: Path) -> dict[str, dict]:
    satirlar, i = _sutunlar(yol, {
        "musteri": MUSTERI, "ad": ("firma adı", "unvan", "müşteri adı", "firma"), "sube": ("şube", "şube adı"),
        "segment": ("segment",), "limit": ("limit", "toplam limit", "kredi limiti"),
        "risk": ("risk", "toplam risk", "nakdi risk", "kullanılan"), "gecikme": ("gecikme günü", "gecikme", "gecikme gün sayısı", "dpd")},
        ("musteri",))
    sonuc = {}
    for r in satirlar:
        m = _al(r, i["musteri"])
        if not m or katla(m).startswith(("toplam", "genel")):
            continue
        sonuc[anahtar(m)] = {"musteri": str(m).strip(), "ad": str(_al(r, i["ad"]) or "").strip(), "sube": str(_al(r, i["sube"]) or "").strip(),
                             "segment": str(_al(r, i["segment"]) or "").strip(),
                             "limit": para(_al(r, i["limit"])) if i["limit"] is not None else None,
                             "risk": para(_al(r, i["risk"])) if i["risk"] is not None else None,
                             "gecikme": int(para(_al(r, i["gecikme"]))) if i["gecikme"] is not None else 0}
    return sonuc


def olay_oku(yol: Path | None, tur_sutun: tuple) -> list[dict]:
    if not yol:
        return []
    satirlar, i = _sutunlar(yol, {"musteri": MUSTERI, "tarih": ("tarih", "kayıt tarihi", "ibraz tarihi", "protesto tarihi", "tebliğ tarihi"),
                                  "tur": tur_sutun, "tutar": ("tutar", "tutarı", "haciz tutarı"),
                                  "durum": ("durum", "ödeme durumu", "kapanış"), "kaynak": ("alacaklı", "haciz koyan", "kurum")},
                            ("musteri", "tarih"))
    return [{"musteri": anahtar(_al(r, i["musteri"])), "tarih": tarih(_al(r, i["tarih"])), "tur": katla(_al(r, i["tur"])),
             "tur_ham": str(_al(r, i["tur"]) or "").strip(), "tutar": para(_al(r, i["tutar"])), "durum": katla(_al(r, i["durum"])),
             "kaynak": str(_al(r, i["kaynak"]) or "").strip()}
            for r in satirlar if _al(r, i["musteri"])]


def ciro_oku(yol: Path | None) -> dict[str, dict[tuple[int, int], Decimal]]:
    if not yol:
        return {}
    satirlar, i = _sutunlar(yol, {"musteri": MUSTERI, "donem": ("dönem", "ay", "tarih"),
                                  "ciro": ("ciro", "net satışlar", "hesap girişleri", "tahsilat", "pos cirosu", "tutar")},
                            ("musteri", "donem", "ciro"))
    sonuc: dict = {}
    for r in satirlar:
        d = donem(_al(r, i["donem"]))
        if _al(r, i["musteri"]) and d:
            k = sonuc.setdefault(anahtar(_al(r, i["musteri"])), {})
            k[d] = k.get(d, SIFIR) + para(_al(r, i["ciro"]))
    return sonuc


# ----------------------------------------------------------------------------
# Puanlama
# ----------------------------------------------------------------------------

def _odenmis(o) -> bool:
    return any(k in o["durum"] for k in ("odendi", "odenmis", "kapandi", "kapali", "temizlendi", "kaldirildi", "fek"))


def puanla(portfoy, cek_senet, haciz, ciro, ref: date, ayar: dict) -> tuple[list[dict], list[dict]]:
    baslangic = ref - timedelta(days=round(ayar["pencere_ay"] * 365 / 12))
    sinyaller: list[dict] = []
    bilinmeyen: set[str] = set()

    def ekle(k, tur, aciklama, puan, tutar=None, t=None):
        sinyaller.append({"anahtar": k, "tur": tur, "aciklama": aciklama, "puan": puan, "tutar": tutar, "tarih": t})

    for k, p in portfoy.items():
        g = p["gecikme"]
        puan = max((pu for esik, pu in ayar["gecikme"] if g >= esik), default=0)
        if puan:
            not_ = " — donuk alacak (Aşama 3) değerlendirmesi" if g > 90 else (" — yakın izleme (Aşama 2) değerlendirmesi" if g > 30 else "")
            ekle(k, "Gecikme", f"{g} gün gecikme{not_}", puan, p["risk"])
        if p["limit"] and p["risk"] is not None:
            oran = p["risk"] / p["limit"] * 100
            if p["risk"] > p["limit"]:
                ekle(k, "Limit aşımı", f"risk {tl(p['risk'])} > limit {tl(p['limit'])} (%{oran:.0f})", ayar["limit"]["asim"], p["risk"] - p["limit"])
            elif oran >= ayar["limit"]["doluluk_esik"]:
                ekle(k, "Limit doluluğu", f"limitin %{oran:.0f}'i kullanılıyor", ayar["limit"]["doluluk"])
        elif p["risk"] and not p["limit"]:
            ekle(k, "Limit aşımı", f"limit tanımsız, risk {tl(p['risk'])}", ayar["limit"]["asim"], p["risk"])

    # Karşılıksız çek / protestolu senet
    for tur, ad in (("karsiliksiz_cek", "Karşılıksız çek"), ("protestolu_senet", "Protestolu senet")):
        kume: dict[str, list] = {}
        for o in cek_senet:
            cek = "cek" in o["tur"] or "karsiliksiz" in o["tur"]
            if (tur == "karsiliksiz_cek") != cek:
                continue
            if o["tarih"] and baslangic <= o["tarih"] <= ref:
                kume.setdefault(o["musteri"], []).append(o)
        for k, liste in kume.items():
            ac = [o for o in liste if not _odenmis(o)]
            od = [o for o in liste if _odenmis(o)]
            puan = min(ayar[tur]["ust_sinir"], len(ac) * ayar[tur]["acik"] + len(od) * ayar[tur]["odenmis"])
            son = max(o["tarih"] for o in liste)
            ekle(k, ad, f"{len(liste)} kayıt ({len(ac)} açık, {len(od)} ödenmiş), son {son:%d.%m.%Y}", puan,
                 sum((o["tutar"] for o in ac), SIFIR), son)
            if k not in portfoy:
                bilinmeyen.add(k)

    # Haciz
    kume = {}
    for o in haciz:
        if o["tarih"] and baslangic <= o["tarih"] <= ref:
            kume.setdefault(o["musteri"], []).append(o)
    kamu_kelime = ("vergi", "sgk", "sosyal guvenlik", "belediye", "gumruk", "kamu", "maliye", "6183")
    for k, liste in kume.items():
        kamu = [o for o in liste if any(w in o["tur"] or w in katla(o["kaynak"]) for w in kamu_kelime)]
        puan = min(ayar["haciz"]["ust_sinir"], len(kamu) * ayar["haciz"]["kamu"] + (len(liste) - len(kamu)) * ayar["haciz"]["ozel"])
        son = max(o["tarih"] for o in liste)
        ekle(k, "Haciz", f"{len(liste)} haciz ihbarnamesi ({len(kamu)} kamu alacağı), son {son:%d.%m.%Y}", puan,
             sum((o["tutar"] for o in liste), SIFIR), son)
        if k not in portfoy:
            bilinmeyen.add(k)

    # Ciro düşüşü
    son_ay = (ref.year, ref.month)
    for k, aylar in ciro.items():
        mevcut = [d for d in aylar if d <= son_ay]
        if not mevcut:
            continue
        sa = max(mevcut)
        son3 = [ay_kaydir(sa, -j) for j in range(3)]
        if not all(d in aylar for d in son3):
            continue
        gy = [ay_kaydir(d, -12) for d in son3]
        on = [ay_kaydir(d, -3) for d in son3]
        if all(d in aylar for d in gy):
            kiyas, ne = gy, "geçen yılın aynı 3 ayına"
        elif all(d in aylar for d in on):
            kiyas, ne = on, "önceki 3 aya"
        else:
            continue
        a, b = sum(aylar[d] for d in son3), sum(aylar[d] for d in kiyas)
        if b <= 0:
            continue
        dusus = (1 - a / b) * 100
        puan = max((pu for esik, pu in ayar["ciro"] if dusus >= esik), default=0)
        if puan:
            ekle(k, "Ciro düşüşü", f"son 3 ay ({son3[-1][1]:02d}-{son3[0][1]:02d}.{son3[0][0]}) {ne} göre %{dusus:.0f} düşüş", puan, b - a)
        if k not in portfoy:
            bilinmeyen.add(k)

    firmalar = []
    for k in list(portfoy) + sorted(bilinmeyen):
        p = portfoy.get(k, {"musteri": k, "ad": "(portföyde yok)", "sube": "", "segment": "", "limit": None, "risk": None, "gecikme": 0})
        s = [x for x in sinyaller if x["anahtar"] == k]
        puan = min(100, sum(x["puan"] for x in s))
        sinif = next(ad for esik, ad in ayar["siniflar"] if puan >= esik)
        turler = sorted({x["tur"] for x in s})
        firmalar.append({**p, "anahtar": k, "puan": puan, "sinif": sinif, "sinyal_sayisi": len(turler),
                         "ozet": "; ".join(f"{x['tur']}: {x['aciklama']}" for x in s), "portfoyde": k in portfoy})
    firmalar.sort(key=lambda f: (-f["puan"], -(f["risk"] or 0), f["musteri"]))
    return firmalar, sinyaller


def calistir(portfoy_yolu: Path, cikti: Path, cek_senet_yolu: Path | None = None, haciz_yolu: Path | None = None,
             ciro_yolu: Path | None = None, ref: date | None = None, ayar_yolu: Path | None = None) -> dict:
    ayar = dict(VARSAYILAN_AYAR)
    if ayar_yolu:
        ayar.update(json.loads(ayar_yolu.read_text(encoding="utf-8")))
    portfoy = portfoy_oku(portfoy_yolu)
    cek_senet = olay_oku(cek_senet_yolu, ("tür", "tip", "kayıt türü", "belge türü"))
    haciz = olay_oku(haciz_yolu, ("tür", "tip", "haciz türü", "alacak türü"))
    ciro = ciro_oku(ciro_yolu)
    if ref is None:
        tarihler = [o["tarih"] for o in cek_senet + haciz if o["tarih"]]
        ref = max(tarihler) if tarihler else date.today()
    firmalar, sinyaller = puanla(portfoy, cek_senet, haciz, ciro, ref, ayar)
    _rapor(firmalar, sinyaller, ref, ayar, cikti)
    return {"firmalar": firmalar, "sinyaller": sinyaller, "ref": ref}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Kritik": "F8C9C6", "Yakın İzleme": "FDE2E1", "İzleme": "FFF4CE", "Normal": "E6F4EA"}
PARA = "#,##0"


def _baslik(ws, satir=1):
    for c in ws[satir]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI


def _f(x):
    return None if x is None else float(x)


def _rapor(firmalar, sinyaller, ref, ayar, cikti):
    wb = Workbook()
    o = wb.active
    o.title = "Özet"
    o.append([f"Erken uyarı sinyalleri — referans tarihi {ref:%d.%m.%Y}"])
    o["A1"].font = Font(bold=True, size=12)
    o.append(["Sınıf", "Firma", "Toplam Risk"])
    _baslik(o, 2)
    for _, ad in ayar["siniflar"]:
        f = [x for x in firmalar if x["sinif"] == ad]
        o.append([ad, len(f), _f(sum((x["risk"] or SIFIR for x in f), SIFIR))])
        o.cell(o.max_row, 1).fill = PatternFill("solid", fgColor=RENK.get(ad, "FFFFFF"))
        o.cell(o.max_row, 3).number_format = PARA
    o.append([])
    o.append(["Sinyal türü", "Firma"])
    _baslik(o, o.max_row)
    for tur in sorted({s["tur"] for s in sinyaller}):
        o.append([tur, len({s["anahtar"] for s in sinyaller if s["tur"] == tur})])
    o.append([])
    o.append(["Not", "Puanlar ayar dosyasındaki örnek değerlerle hesaplanmıştır; kurum kredi politikasına göre ayarlayın. "
                     "Araç karar vermez; sınıflandırma ve aşama (TFRS 9 / BDDK karşılık yönetmeliği) kararı yetkili birimindir."])
    o.column_dimensions["A"].width = 22
    o.column_dimensions["B"].width = 14
    o.column_dimensions["C"].width = 18

    f = wb.create_sheet("Firma Listesi")
    f.append(["Sıra", "Müşteri No", "Firma", "Şube", "Segment", "Limit", "Risk", "Gecikme Günü", "Puan", "Sınıf", "Sinyal Türü Sayısı",
              "Sinyaller", "İzleme Kararı", "Açıklama"])
    _baslik(f)
    for n, x in enumerate(firmalar, 1):
        f.append([n, x["musteri"], x["ad"], x["sube"], x["segment"], _f(x["limit"]), _f(x["risk"]), x["gecikme"], x["puan"], x["sinif"],
                  x["sinyal_sayisi"], x["ozet"], None, None])
        f.cell(f.max_row, 10).fill = PatternFill("solid", fgColor=RENK.get(x["sinif"], "FFFFFF"))
        for c in (6, 7):
            f.cell(f.max_row, c).number_format = PARA
    for j, w in enumerate((6, 12, 28, 14, 10, 13, 13, 9, 7, 13, 9, 90, 18, 30), 1):
        f.column_dimensions[get_column_letter(j)].width = w
    f.freeze_panes = "D2"
    f.auto_filter.ref = f.dimensions

    s = wb.create_sheet("Sinyal Detayı")
    s.append(["Müşteri No", "Firma", "Sinyal", "Açıklama", "Puan", "Tutar", "Son Tarih"])
    _baslik(s)
    ad = {x["anahtar"]: x for x in firmalar}
    for x in sorted(sinyaller, key=lambda x: (-ad[x["anahtar"]]["puan"], x["anahtar"], -x["puan"])):
        s.append([ad[x["anahtar"]]["musteri"], ad[x["anahtar"]]["ad"], x["tur"], x["aciklama"], x["puan"], _f(x["tutar"]), x["tarih"]])
        s.cell(s.max_row, 6).number_format = PARA
        s.cell(s.max_row, 7).number_format = "DD.MM.YYYY"
    for j, w in enumerate((12, 28, 18, 70, 7, 14, 12), 1):
        s.column_dimensions[get_column_letter(j)].width = w
    s.freeze_panes = "C2"
    s.auto_filter.ref = s.dimensions

    b = wb.create_sheet("Puan Tablosu")
    b.append(["Sinyal", "Kural", "Puan"])
    _baslik(b)
    for esik, pu in ayar["gecikme"]:
        b.append(["Gecikme", f"≥ {esik} gün", pu])
    for tur, ad_ in (("karsiliksiz_cek", "Karşılıksız çek"), ("protestolu_senet", "Protestolu senet")):
        b.append([ad_, f"açık kayıt başına (son {ayar['pencere_ay']} ay)", ayar[tur]["acik"]])
        b.append([ad_, "ödenmiş kayıt başına", ayar[tur]["odenmis"]])
        b.append([ad_, "üst sınır", ayar[tur]["ust_sinir"]])
    b.append(["Haciz", "kamu alacağı (vergi, SGK…) başına", ayar["haciz"]["kamu"]])
    b.append(["Haciz", "özel alacaklı başına", ayar["haciz"]["ozel"]])
    b.append(["Haciz", "üst sınır", ayar["haciz"]["ust_sinir"]])
    b.append(["Limit", "risk > limit", ayar["limit"]["asim"]])
    b.append(["Limit", f"doluluk ≥ %{ayar['limit']['doluluk_esik']}", ayar["limit"]["doluluk"]])
    for esik, pu in ayar["ciro"]:
        b.append(["Ciro düşüşü", f"≥ %{esik}", pu])
    for esik, ad_ in ayar["siniflar"]:
        b.append(["Sınıf", f"puan ≥ {esik} → {ad_}", None])
    b.append(["Toplam", "sinyal puanları toplanır, en çok 100", None])
    b.column_dimensions["A"].width = 18
    b.column_dimensions["B"].width = 40
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ornek = BURASI / "ornek_veri"
    ap = argparse.ArgumentParser(description="Kredi müşterilerinin erken uyarı sinyallerini birleştirip firma bazında risk puanı üretir.")
    ap.add_argument("--portfoy", type=Path, default=ornek / "portfoy.csv", help="Müşteri No, Firma Adı, Şube, Limit, Risk, Gecikme Günü")
    ap.add_argument("--cek-senet", type=Path, help="Müşteri No, Tarih, Tür (Karşılıksız Çek / Protestolu Senet), Tutar, Durum")
    ap.add_argument("--haciz", type=Path, help="Müşteri No, Tarih, Alacaklı, Tür, Tutar")
    ap.add_argument("--ciro", type=Path, help="Müşteri No, Dönem (2026-09), Ciro")
    ap.add_argument("--tarih", help="Referans tarihi (gg.aa.yyyy); varsayılan: verideki en son tarih")
    ap.add_argument("--ayar", type=Path, help="Puan ve eşik ayarları (JSON)")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "erken_uyari_listesi.xlsx")
    a = ap.parse_args(argv)
    if a.portfoy == ornek / "portfoy.csv":
        a.cek_senet = a.cek_senet or ornek / "kkb_cek_senet.csv"
        a.haciz = a.haciz or ornek / "haciz.csv"
        a.ciro = a.ciro or ornek / "ciro.csv"
        a.tarih = a.tarih or "30.09.2026"
    ref = tarih(a.tarih) if a.tarih else None
    if a.tarih and ref is None:
        raise SystemExit(f"Tarih anlaşılamadı: {a.tarih}")
    s = calistir(a.portfoy, a.cikti, a.cek_senet, a.haciz, a.ciro, ref, a.ayar)
    from collections import Counter
    say = Counter(f["sinif"] for f in s["firmalar"])
    print(f"[OK] {len(s['firmalar'])} firma · " + " · ".join(f"{ad}: {say.get(ad, 0)}" for _, ad in VARSAYILAN_AYAR["siniflar"]))
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
