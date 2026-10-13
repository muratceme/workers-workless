import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)   # Pazartesi
A, Y, N = "1 – Acil", "2 – Yüksek", "3 – Normal"


def h(no, tutar, kalan, pb="TRY", yenileme=False, py="P", urun="Vadeli", gorusme=None):
    return {"no": no, "ad": "Ad", "urun": urun, "pb": pb, "tutar": D(str(tutar)), "vade": BUGUN + timedelta(days=kalan), "faiz": None, "yenileme": yenileme, "gorusme": gorusme, "py": py}


class DegerlendirTesti(unittest.TestCase):
    def test_pencere_ve_oncelik(self):
        l = [h("bugün", 10, 0), h("üç", 10, 3), h("dört", 10, 4), h("onbeş", 10, 15), h("onaltı", 10, 16), h("geçmiş", 10, -2), h("büyük", 1000000, 8), h("altı", "999999.99", 8), h("acil_büyük", 5000000, 1)]
        r = main.degerlendir(l, BUGUN)
        self.assertEqual([(o["no"], o["oncelik"], o["kalan"], o["gerekce"]) for o in r["musteriler"]],
                         [("acil_büyük", A, 1, "Vadeye 1 gün kaldı"), ("geçmiş", A, -2, "Vade 2 gün önce doldu"), ("bugün", A, 0, "Vade bugün"), ("üç", A, 3, "Vadeye 3 gün kaldı"), ("büyük", Y, 8, "Vadesi dolan tutar 1,000,000.00 TL (eşik 1,000,000)"),
                          ("altı", N, 8, ""), ("dört", N, 4, ""), ("onbeş", N, 15, "")])
        self.assertEqual((r["say"], r["toplam"], len(r["hesaplar"]), len(r["tum"]), [h_["no"] for h_ in r["hesaplar"]][:2]), ({A: 4, Y: 1, N: 3}, D("7000049.99"), 8, 9, ["geçmiş", "bugün"]))
        self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]], [("Orta", "Vadesi geçmiş kayıt", "geçmiş · Ad")])
        o = lambda **kw: {x["no"]: x["oncelik"] for x in main.degerlendir(l, BUGUN, **kw)["musteriler"]}   # noqa: E731
        self.assertEqual((o(gun=16)["onaltı"], o(acil=4)["dört"], o(acil=0)["üç"], o(buyuk=D(10))["dört"], "onbeş" in o(gun=14)), (N, A, N, Y, False))
        with self.assertRaises(ValueError):
            main.degerlendir([], BUGUN)

    def test_musteri_ve_doviz(self):
        l = [h("M1", 750000, 2), h("m1", 20000, 8, "USD", True, urun="Vadeli", gorusme=date(2026, 9, 1)), h("M1", 100, 40), h("M2", 15000, 5, "EUR", py="R"), h("M3", 10000, 7, "GBP", py="R"), h("M3", 500, 9, py="R", urun="Fon", gorusme=date(2026, 8, 1)),
             h("M3", 500, 10, py="R", urun="Fon", gorusme=date(2026, 9, 9)), h("M4", 100, 6, py="r")]
        r = main.degerlendir(l, BUGUN, {"USD": D("42.50"), "EUR": D(46)})
        self.assertEqual([(o["no"], o["oncelik"], o["hesap"], o["toplam"], o["en_yakin"], o["kalan"], o["yenilemesiz"], o["gorusme"], o["urunler"], o["py"]) for o in r["musteriler"]],
                         [("M1", A, 2, D(1600000), date(2026, 10, 14), 2, 1, date(2026, 9, 1), "Vadeli (TRY), Vadeli (USD)", "P"), ("M2", N, 1, D(690000), date(2026, 10, 17), 5, 1, None, "Vadeli (EUR)", "R"),
                          ("M3", N, 3, D(1000), date(2026, 10, 19), 7, 3, date(2026, 9, 9), "Vadeli (GBP), Fon (TRY)", "R"), ("M4", N, 1, D(100), date(2026, 10, 18), 6, 1, None, "Vadeli (TRY)", "r")])
        self.assertEqual(r["para_birimleri"], [{"pb": "EUR", "hesap": 1, "tutar": D(15000), "tl": D(690000)}, {"pb": "GBP", "hesap": 1, "tutar": D(10000), "tl": None}, {"pb": "TRY", "hesap": 4, "tutar": D(751100), "tl": D(751100)},
                                               {"pb": "USD", "hesap": 1, "tutar": D(20000), "tl": D(850000)}])
        self.assertEqual(r["yoneticiler"], [{"py": "P", "musteri": 1, "toplam": D(1600000), "acil": 1}, {"py": "R", "musteri": 3, "toplam": D(691100), "acil": 0}])
        self.assertEqual([(t["vade"].day, t["hesap"], t["tl"]) for t in r["takvim"]], [(14, 1, D(750000)), (17, 1, D(690000)), (18, 1, D(100)), (19, 1, D(0)), (20, 1, D(850000)), (21, 1, D(500)), (22, 1, D(500))])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Bilgi", "Vade hafta sonuna denk geliyor", "M2 · Ad"), ("Bilgi", "Vade hafta sonuna denk geliyor", "M4 · Ad"), ("Yüksek", "Kur yok", "M3 · Ad")])
        self.assertEqual(main.degerlendir([h("x", 5, 2, "USD")], BUGUN)["musteriler"][0]["toplam"], D(0))


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "hesaplar.csv", cikti, BUGUN, ORNEK / "kurlar.csv")
            self.assertEqual([(o["no"], o["ad"], o["oncelik"], o["hesap"], o["toplam"], o["en_yakin"], o["kalan"], o["yenilemesiz"]) for o in r["musteriler"]],
                             [("1001", "Ahmet Yılmaz", A, 2, D(1600000), date(2026, 10, 14), 2, 1), ("1005", "Ali Şahin", A, 1, D(500000), date(2026, 10, 9), -3, 1), ("1002", "Fatma Demir", Y, 2, D(3000000), date(2026, 10, 22), 10, 2),
                              ("1004", "Ayşe Çelik", N, 1, D(690000), date(2026, 10, 17), 5, 1), ("1003", "Mehmet Kaya", N, 1, D(300000), date(2026, 10, 26), 14, 1), ("1007", "Hasan Koç", N, 1, D(0), date(2026, 10, 19), 7, 1)])
            self.assertEqual((r["toplam"], r["say"], len(r["hesaplar"]), len(r["tum"])), (D(6090000), {A: 2, Y: 1, N: 3}, 8, 9))
            self.assertEqual(r["para_birimleri"], [{"pb": "EUR", "hesap": 1, "tutar": D(15000), "tl": D(690000)}, {"pb": "GBP", "hesap": 1, "tutar": D(10000), "tl": None}, {"pb": "TRY", "hesap": 5, "tutar": D(4550000), "tl": D(4550000)},
                                                   {"pb": "USD", "hesap": 1, "tutar": D(20000), "tl": D(850000)}])
            self.assertEqual(r["yoneticiler"], [{"py": "Okan Er", "musteri": 4, "toplam": D(1490000), "acil": 1}, {"py": "Zeynep Ak", "musteri": 2, "toplam": D(4600000), "acil": 1}])
            self.assertEqual([(t["vade"], t["tl"]) for t in r["takvim"]][:3], [(date(2026, 10, 9), D(500000)), (date(2026, 10, 14), D(750000)), (date(2026, 10, 17), D(690000))])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Yüksek", "Kur yok", "1007 · Hasan Koç"), ("Orta", "Vadesi geçmiş kayıt", "1005 · Ali Şahin"), ("Bilgi", "Vade hafta sonuna denk geliyor", "1004 · Ayşe Çelik")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Görüşme Listesi", "Hesaplar", "Özet", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, ky = Path(tmp) / "h.csv", Path(tmp) / "k.csv"
            y.write_text("No;Ad Soyad;Hesap Türü;Döviz;Bakiye;Vade;Faiz;Temdit;Son Temas;Portföy\n7;Ali;;tl;1.250.000,50;15.10.2026;42,5;EVET;01.10.2026;\n8;Ece;Fon;Dolar;5.000;2026-10-20;;;;Can\n9;X;;;0;15.10.2026;;;;\n10;X;;;5;x;;;;\n"
                         ";X;;;5;15.10.2026;;;;\n", encoding="utf-8")
            ky.write_text("Döviz Cinsi;Alış Kuru\nDolar;42,5\n;3\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["urun"], x["pb"], x["tutar"], x["vade"], x["faiz"], x["yenileme"], x["gorusme"], x["py"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("7", "(ürün yok)", "TRY", D("1250000.50"), date(2026, 10, 15), D("42.5"), True, date(2026, 10, 1), "(atanmamış)"), ("8", "Fon", "USD", D(5000), date(2026, 10, 20), None, False, None, "Can")],
                              [("Hesap okunamadı", "9 · X"), ("Hesap okunamadı", "10 · X")]))
            self.assertEqual((main.oku_kurlar(ky), [main.pb_kod(x) for x in ("", "TL", "usd", "Euro", "Sterlin", "chf")]), ({"USD": D("42.5")}, ["TRY", "TRY", "USD", "EUR", "GBP", "CHF"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--hesaplar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--kurlar", str(ky), "--tarih", "12.10.2026", "--gun", "30", "--acil", "5", "--buyuk", "2.500.000"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--gun", "-1"], ["--acil", "20"], ["--buyuk", "0"], ["--buyuk", "x"])], [2] * 5)
            self.assertEqual(main.main(arg + ["--kurlar", str(Path(tmp) / "yok.csv")]), 1)
            ky.write_text("Para Birimi;Kur\nUSD;0\n", encoding="utf-8")
            self.assertEqual(main.main(arg + ["--kurlar", str(ky)]), 1)
            for icerik in ("Müşteri No;Tutar\nA;1\n", "Müşteri No;Tutar;Vade\n;1;15.10.2026\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
