"""
Vadesi Dolan Mevduat Takibi — Gerisi Bende kod bloğu
Şube Bankacılığı › Bireysel Portföy Yöneticisi

Vadesi yaklaşan mevduat ve yatırım ürünlerini müşteri bazında listeler; görüşme önceliği önerir.
  Kalan gün          = vade tarihi − durum tarihi
  Pencere            = kalan günü 0 ile --gun (15) arasında olan hesaplar; vadesi geçmiş kayıtlar da listeye girer ve uyarılır
  TL karşılığı       = tutar × kur   (TRY için 1; kurlar --kurlar dosyasından)
  Müşteri toplamı    = müşterinin penceredeki hesaplarının TL karşılığı toplamı
  Öncelik            = 1 – Acil    : en yakın vadeye --acil (3) gün veya az kalmış ya da vade geçmiş
                       2 – Yüksek  : müşteri toplamı --buyuk (1.000.000 TL) ve üstünde
                       3 – Normal  : diğerleri
  Sıra               = öncelik → müşteri toplamı (büyükten küçüğe) → en yakın vade
Öncelik kuralı bir çalışma varsayımıdır; eşikler komut satırından değişir. Kod ürün önermez ve faiz teklifi hesaplamaz.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
SUTUN = {"no": ("musteri no", "musteri numarasi", "no"), "ad": ("musteri", "musteri adi", "ad soyad"), "urun": ("urun", "urun turu", "hesap turu"), "pb": ("para birimi", "doviz", "doviz cinsi", "pb"), "tutar": ("tutar", "bakiye", "anapara"),
         "vade": ("vade tarihi", "vade sonu", "vade"), "faiz": ("faiz orani", "faiz", "oran"), "yenileme": ("otomatik yenileme", "temdit", "otomatik temdit"), "gorusme": ("son gorusme tarihi", "son gorusme", "son temas"),
         "py": ("portfoy yoneticisi", "musteri temsilcisi", "portfoy")}
KUR_SUTUN = {"pb": ("para birimi", "doviz", "doviz cinsi", "pb"), "kur": ("kur", "tl karsiligi", "alis kuru")}
EVET = {"evet", "e", "var", "x", "true", "1"}
ONCELIK_RENK = {"1 – Acil": KIRMIZI, "2 – Yüksek": SARI, "3 – Normal": None}
Z = Decimal(0)


def olcu(x) -> Decimal | None:
    """Kur ve oran okur: virgül ya da tek nokta ondalık ayırıcıdır (42,50 = 42.50)."""
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x or ""))
    try:
        return Decimal(s.replace(".", "").replace(",", ".") if "," in s else s) if s else None
    except InvalidOperation:
        return None


def pb_kod(x) -> str:
    k = katla(x).upper()
    return {"": "TRY", "TL": "TRY", "TURK LIRASI": "TRY", "DOLAR": "USD", "ABD DOLARI": "USD", "EURO": "EUR", "AVRO": "EUR", "STERLIN": "GBP"}.get(k, k)


def oku_kurlar(yol: Path) -> dict[str, Decimal]:
    kurlar = {}
    for r in kayitlar(yol, KUR_SUTUN, ("pb", "kur")):
        if not metin(r["pb"]):
            continue
        kur = olcu(r["kur"])
        if kur is None or kur <= 0:
            raise ValueError(f"{yol.name} satır {r['_satir']}: kur '{metin(r['kur'])}' pozitif bir sayı olmalı.")
        kurlar[pb_kod(r["pb"])] = kur
    return kurlar


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("no", "tutar", "vade")):
        no = metin(r["no"])
        if not no:
            continue
        kim = f"{no} · {metin(r['ad'])}".strip(" ·")
        tutar, vade = para(r["tutar"]), tarih(r["vade"])
        if tutar is None or tutar <= 0 or vade is None:
            uy.append(uyari("Yüksek", "Hesap okunamadı", kim, f"Satır {r['_satir']}: tutar '{metin(r['tutar'])}', vade '{metin(r['vade'])}'. Satır alınmadı."))
            continue
        l.append({"no": no, "ad": metin(r["ad"]), "urun": metin(r["urun"]) or "(ürün yok)", "pb": pb_kod(r["pb"]), "tutar": tutar, "vade": vade, "faiz": olcu(r["faiz"]), "yenileme": katla(r["yenileme"]) in EVET,
                  "gorusme": tarih(r["gorusme"]), "py": metin(r["py"]) or "(atanmamış)"})
    return l, uy


def degerlendir(hesaplar: list[dict], bugun: date, kurlar: dict[str, Decimal] | None = None, gun: int = 15, acil: int = 3, buyuk: Decimal = Decimal(1000000)) -> dict:
    if not hesaplar:
        raise ValueError("Değerlendirilecek hesap yok.")
    kurlar = {"TRY": Decimal(1)} | (kurlar or {})
    uy, musteriler = [], {}
    for h in hesaplar:
        kim = f"{h['no']} · {h['ad']}".strip(" ·")
        h["kalan"] = (h["vade"] - bugun).days
        h["tl"] = yuvarla(h["tutar"] * kurlar[h["pb"]]) if h["pb"] in kurlar else None
        h["pencere"] = h["kalan"] <= gun
        if not h["pencere"]:
            continue
        if h["tl"] is None:
            uy.append(uyari("Yüksek", "Kur yok", kim, f"{h['tutar']:,.2f} {h['pb']} için kur verilmedi; TL karşılığı hesaplanamadı ve müşteri toplamına katılmadı (--kurlar)."))
        if h["kalan"] < 0:
            uy.append(uyari("Orta", "Vadesi geçmiş kayıt", kim, f"{h['urun']}: vade {h['vade']:%d.%m.%Y} ({-h['kalan']} gün önce). Hesap yenilendiyse ya da kapandıysa liste güncel değil; değilse müşteriyle hemen görüşün."))
        elif h["vade"].weekday() >= 5:
            uy.append(uyari("Bilgi", "Vade hafta sonuna denk geliyor", kim, f"{h['urun']}: vade {h['vade']:%d.%m.%Y} hafta sonu. İşlem günü kaymış olabilir; görüşmeyi öne alın. Resmî tatiller ayrıca kontrol edilmelidir."))
        m = musteriler.setdefault(katla(h["no"]), {"no": h["no"], "ad": h["ad"], "py": h["py"], "hesaplar": []})
        m["hesaplar"].append(h)
    liste = []
    for m in musteriler.values():
        hs = sorted(m["hesaplar"], key=lambda h: h["vade"])
        o = {"no": m["no"], "ad": m["ad"], "py": m["py"], "hesap": len(hs), "toplam": sum((h["tl"] or Z for h in hs), Z), "en_yakin": hs[0]["vade"], "kalan": hs[0]["kalan"], "yenilemesiz": sum(1 for h in hs if not h["yenileme"]),
             "gorusme": max((h["gorusme"] for h in hs if h["gorusme"]), default=None), "urunler": ", ".join(dict.fromkeys(f"{h['urun']} ({h['pb']})" for h in hs))}
        if o["kalan"] <= acil:
            o["oncelik"], o["gerekce"] = "1 – Acil", f"Vade {-o['kalan']} gün önce doldu" if o["kalan"] < 0 else "Vade bugün" if o["kalan"] == 0 else f"Vadeye {o['kalan']} gün kaldı"
        elif o["toplam"] >= buyuk:
            o["oncelik"], o["gerekce"] = "2 – Yüksek", f"Vadesi dolan tutar {o['toplam']:,.2f} TL (eşik {buyuk:,.0f})"
        else:
            o["oncelik"], o["gerekce"] = "3 – Normal", ""
        liste.append(o)
    liste.sort(key=lambda o: (o["oncelik"], -o["toplam"], o["en_yakin"], katla(o["no"])))
    pencere = sorted((h for h in hesaplar if h["pencere"]), key=lambda h: (h["vade"], katla(h["no"])))

    def grupla(anahtar, liste_) -> list[tuple]:
        g = {}
        for x in liste_:
            g.setdefault(anahtar(x), []).append(x)
        return sorted(g.items())

    return {"musteriler": liste, "hesaplar": pencere, "tum": sorted(hesaplar, key=lambda h: (h["vade"], katla(h["no"]))), "toplam": sum((o["toplam"] for o in liste), Z), "say": {k: sum(1 for o in liste if o["oncelik"] == k) for k in ONCELIK_RENK},
            "para_birimleri": [{"pb": pb, "hesap": len(g), "tutar": sum((h["tutar"] for h in g), Z), "tl": sum((h["tl"] for h in g), Z) if all(h["tl"] is not None for h in g) else None} for pb, g in grupla(lambda h: h["pb"], pencere)],
            "yoneticiler": [{"py": g[0]["py"], "musteri": len(g), "toplam": sum((o["toplam"] for o in g), Z), "acil": sum(1 for o in g if o["oncelik"] == "1 – Acil")} for _, g in grupla(lambda o: katla(o["py"]), liste)],
            "takvim": [{"vade": v, "hesap": len(g), "tl": sum((h["tl"] or Z for h in g), Z)} for v, g in grupla(lambda h: h["vade"], pencere)], "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, bugun: date, gun: int) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Görüşme Listesi"
    _baslik(ws, ["Öncelik", "Müşteri No", "Müşteri", "Vadesi Dolan (TL)", "Hesap", "En Yakın Vade", "Kalan Gün", "Otomatik Yenilemesiz", "Ürünler", "Son Görüşme", "Portföy Yöneticisi", "Gerekçe", "Görüşme Notu (siz doldurun)"],
            (12, 12, 26, 18, 7, 13, 9, 11, 36, 12, 22, 40, 34))
    for o in r["musteriler"]:
        _satir(ws, [o["oncelik"], o["no"], o["ad"], o["toplam"], o["hesap"], o["en_yakin"], o["kalan"], o["yenilemesiz"], o["urunler"], o["gorusme"], o["py"], o["gerekce"], None], para_s=(4,), tarih_s=(6, 10),
               dolgu={1: ONCELIK_RENK[o["oncelik"]]}, sar=(9, 12))
    ws = wb.create_sheet("Hesaplar")
    _baslik(ws, ["Vade", "Kalan Gün", "Müşteri No", "Müşteri", "Ürün", "Para Birimi", "Tutar", "TL Karşılığı", "Faiz %", "Otomatik Yenileme", "Portföy Yöneticisi"], (12, 9, 12, 26, 24, 9, 16, 16, 9, 11, 22))
    for h in r["hesaplar"]:
        _satir(ws, [h["vade"], h["kalan"], h["no"], h["ad"], h["urun"], h["pb"], h["tutar"], h["tl"], h["faiz"], "Evet" if h["yenileme"] else "Hayır", h["py"]], para_s=(7, 8, 9), tarih_s=(1,), dolgu={2: KIRMIZI if h["kalan"] < 0 else None})
    ws = wb.create_sheet("Özet")
    _baslik(ws, ["Gösterge", "Değer"], (40, 18))
    ws.append([f"Durum tarihi {bugun:%d.%m.%Y} · pencere {gun} gün"])
    _satir(ws, ["Görüşülecek müşteri", len(r["musteriler"])])
    for k, n in r["say"].items():
        _satir(ws, [k, n], dolgu={1: ONCELIK_RENK[k]})
    _satir(ws, ["Vadesi dolan toplam (TL)", r["toplam"]], para_s=(2,))
    ws.append([])
    _baslik(ws, ["Para Birimi", "Hesap", "Tutar", "TL Karşılığı"], dondur="A2")
    for p in r["para_birimleri"]:
        _satir(ws, [p["pb"], p["hesap"], p["tutar"], p["tl"]], para_s=(3, 4))
    ws.append([])
    _baslik(ws, ["Portföy Yöneticisi", "Müşteri", "Vadesi Dolan (TL)", "Acil"], dondur="A2")
    for y in r["yoneticiler"]:
        _satir(ws, [y["py"], y["musteri"], y["toplam"], y["acil"]], para_s=(3,))
    ws.append([])
    _baslik(ws, ["Vade Tarihi", "Hesap", "TL Karşılığı"], dondur="A2")
    for t in r["takvim"]:
        _satir(ws, [t["vade"], t["hesap"], t["tl"]], para_s=(3,), tarih_s=(1,))
    ws.column_dimensions["C"].width, ws.column_dimensions["D"].width = 18, 18
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, kur_yol: Path | None = None, gun: int = 15, acil: int = 3, buyuk: Decimal = Decimal(1000000)) -> dict:
    hesaplar, uy = oku(yol)
    r = degerlendir(hesaplar, bugun, oku_kurlar(kur_yol) if kur_yol else None, gun, acil, buyuk)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, bugun, gun)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Vadesi yaklaşan mevduat ve yatırım ürünlerini müşteri bazında listeler, görüşme önceliği önerir.")
    p.add_argument("--hesaplar", type=Path, help="Müşteri No, Müşteri, Ürün, Para Birimi, Tutar, Vade Tarihi, Faiz Oranı, Otomatik Yenileme, Son Görüşme Tarihi, Portföy Yöneticisi")
    p.add_argument("--kurlar", type=Path, help="Para Birimi, Kur (TL karşılığı için; dövizli hesap varsa)")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--gun", type=int, default=15, help="Vadesine kaç gün kalan hesaplar listelensin (varsayılan 15)")
    p.add_argument("--acil", type=int, default=3, help="Vadeye bu kadar gün veya az kalan müşteri 'Acil' sayılır (varsayılan 3)")
    p.add_argument("--buyuk", default="1000000", help="Müşteri toplamı bu TL tutarı ve üstündeyse 'Yüksek' öncelik (varsayılan 1.000.000)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "vadesi_dolan_mevduat_takibi.xlsx")
    a = p.parse_args(argv)
    ornek = a.hesaplar is None
    yol = ORNEK / "hesaplar.csv" if ornek else a.hesaplar
    kur_yol = ORNEK / "kurlar.csv" if ornek and a.kurlar is None else a.kurlar
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    buyuk = para(a.buyuk)
    if bugun is None or not 0 <= a.acil <= a.gun <= 365 or buyuk is None or buyuk <= 0:
        print("[X] --tarih GG.AA.YYYY; --gun 0–365; --acil sıfır veya pozitif ve --gun değerinden büyük değil; --buyuk pozitif olmalı")
        return 2
    for y in (yol, kur_yol):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(yol, a.cikti, bugun, kur_yol, a.gun, a.acil, buyuk)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    s = r["say"]
    print(f"[OK] {len(r['hesaplar'])} hesap, {len(r['musteriler'])} müşteri · vadesi dolan {r['toplam']:,.2f} TL · Acil {s['1 – Acil']}, Yüksek {s['2 – Yüksek']}, Normal {s['3 – Normal']}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
