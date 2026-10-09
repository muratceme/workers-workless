import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "e.xlsx"
        cls.s = main.calistir(ORNEK / "yetkinlik_matrisi.csv", ORNEK / "degerlendirmeler.csv", cls.cikti, ORNEK / "performans.csv",
                              ORNEK / "talepler.csv")
        cls.i = {(x.sicil, x.g.yetkinlik): x for x in cls.s["ihtiyaclar"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_puan_ve_oncelik(self):
        x = self.i[("4013", "e-Dönüşüm uygulamaları")]        # açık 2 × 3 + düşük performans + talep
        self.assertEqual((x.acik, x.puan, x.oncelik, x.talep), (2, 8, 1, True))
        x = self.i[("4006", "Ürün bilgisi")]                  # yönetici puanı yok → öz değerlendirme
        self.assertEqual((x.mevcut, x.oncelik, x.notlar), (2, 1, ["Yalnız öz değerlendirme"]))
        self.assertEqual(self.i[("4021", "İleri Excel")].oncelik, 3)
        self.assertIsNone(self.i[("4004", "Müzakere")].oncelik)
        self.assertEqual(len(self.s["ihtiyaclar"]), 44)

    def test_oncelik_ve_plan(self):
        o = self.s["oncelik"][0]
        self.assertEqual((o["departman"], o["yetkinlik"], o["acikli"], o["puan"]), ("Satış", "Ürün bilgisi", 5, 23))
        p = {x["egitim"]: x for x in self.s["plan"]}
        self.assertEqual((p["Ürün Akademisi"]["kisi"], p["Ürün Akademisi"]["yontem"][:4]), (5, "Grup"))
        self.assertTrue(p["İleri Excel"]["yontem"].startswith("Bireysel"))
        self.assertEqual(p["İleri Excel"]["departmanlar"], "Muhasebe, Üretim")
        oz = {x["departman"]: x for x in self.s["ozet"]}
        self.assertEqual((oz["Satış"]["degerlendirme"], oz["Satış"]["karsilanan"], oz["Satış"]["o1"]), (23, 8, 3))

    def test_algi_talep_uyari(self):
        self.assertEqual(sorted(x["i"].sicil for x in self.s["farklar"]), ["4001", "4003", "4011", "4021", "4023"])
        t = {x["sicil"]: x["durum"] for x in self.s["talepler"]}
        self.assertEqual(t["4002"], "Talep ve ihtiyaç örtüşüyor")
        self.assertEqual(t["4004"], "Matristeki eğitimlerle eşleşmedi")
        u = {(x["tur"], x["kim"]) for x in self.s["uyarilar"]}
        self.assertTrue({("Değerlendirilmedi", "4004"), ("Matriste olmayan yetkinlik", "4023")} <= u)

    def test_excel(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Öncelik Listesi", "Eğitim Planı", "Bireysel İhtiyaçlar", "Departman Özeti", "Algı Farkları", "Talepler", "Uyarılar"])
        self.assertEqual(wb["Öncelik Listesi"].max_row, 12)
        self.assertEqual(wb["Bireysel İhtiyaçlar"]["N2"].value, 1)


class KuralTesti(unittest.TestCase):
    def test_min_grup_ve_eksik_dosyalar(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            s = main.calistir(ORNEK / "yetkinlik_matrisi.csv", ORNEK / "degerlendirmeler.csv", t / "o.xlsx", min_grup=3)
            p = {x["egitim"]: x for x in s["plan"]}
            self.assertTrue(p["İleri Excel"]["yontem"].startswith("Grup"))
            self.assertEqual(s["talepler"], [])
            (t / "d.csv").write_text("Sicil No;Pozisyon;Yetkinlik;Yönetici Puanı;Öz Değerlendirme\n1;Kasiyer;Hız;3;3\n2;Ekip Lideri;Planlama;2,5;9\n",
                                     encoding="utf-8")
            (t / "p.csv").write_text("Sicil No;Performans\n2;2\n", encoding="utf-8")
            s = main.calistir(ORNEK / "yetkinlik_matrisi.csv", t / "d.csv", t / "o.xlsx", t / "p.csv")
            x = next(i for i in s["ihtiyaclar"] if i.sicil == "2")
            self.assertEqual((x.acik, x.puan, x.oz), (1.5, 5.5, None))     # öz 9 geçersiz; performans 2 = düşük
            turler = {u["tur"] for u in s["uyarilar"]}
            self.assertTrue({"Pozisyon matriste yok", "Değerlendirilmedi"} <= turler)

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--matris", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
