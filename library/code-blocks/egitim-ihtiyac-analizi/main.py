"""
Eğitim İhtiyaç Analizi — Workers / Workless kod bloğu
İnsan Kaynakları › Eğitim ve Gelişim Uzmanı

Pozisyon yetkinlik matrisini, yönetici ve öz değerlendirmelerini, performans sonuçlarını ve eğitim taleplerini
birleştirerek departman bazında önceliklendirilmiş eğitim ihtiyacı çıkarır:
  - Mevcut seviye = yönetici puanı (yoksa öz değerlendirme, işaretlenir). Açık = gerekli seviye − mevcut seviye.
  - İhtiyaç puanı = açık × kritiklik ağırlığı (Yüksek 3, Orta 2, Düşük 1) + 1 (performans C/D veya 1–2)
    + 1 (çalışan bu eğitimi talep etmiş). Öncelik: puan ≥ 6 → 1, 3–5 → 2, 1–2 → 3.
  - Algı farkı: öz değerlendirme yönetici puanından 2+ yüksekse "kör nokta", 2+ düşükse "fark edilmemiş güç".
  - Eğitim planı: aynı önerilen eğitime ihtiyacı olanlar gruplanır; katılımcı sayısı --min-grup ve üstüyse
    grup eğitimi, altıysa bireysel yöntem (e-öğrenme, mentorluk, iş başında) önerilir.
  - Matriste olmayan yetkinlik, değerlendirilmemiş yetkinlik ve ihtiyaçla eşleşmeyen talepler ayrıca listelenir.
Rapor: öncelik listesi, eğitim planı, bireysel ihtiyaçlar, departman özeti, algı farkları, talepler, uyarılar.
İnternete bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: 12 çalışan, 3 pozisyon
    python main.py --matris matris.xlsx --degerlendirmeler d.xlsx --performans p.xlsx --talepler t.xlsx
    python main.py --matris m.csv --degerlendirmeler d.csv --min-grup 8
"""
from __future__ import annotations

import argparse
import csv
import re
import statistics
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
KRITIKLIK = {"yuksek": 3, "orta": 2, "dusuk": 1, "3": 3, "2": 2, "1": 1}
KRITIK_AD = {3: "Yüksek", 2: "Orta", 1: "Düşük"}
DUSUK_PERF = {"c", "d", "1", "2"}

MATRIS_SUTUNLARI = {"pozisyon": ("pozisyon", "unvan"), "yetkinlik": ("yetkinlik", "beceri"), "gerekli": ("gerekli seviye", "beklenen seviye", "hedef seviye"),
                    "kritiklik": ("kritiklik", "onem"), "egitim": ("onerilen egitim", "egitim", "egitim programi")}
DEGER_SUTUNLARI = {"sicil": ("sicil no", "sicil"), "ad": ("ad soyad", "ad"), "departman": ("departman", "birim"), "pozisyon": ("pozisyon", "unvan"),
                   "yetkinlik": ("yetkinlik", "beceri"), "yonetici": ("yonetici puani", "yonetici degerlendirmesi", "yonetici"),
                   "oz": ("oz degerlendirme", "oz puan", "calisan puani")}
PERF_SUTUNLARI = {"sicil": ("sicil no", "sicil"), "performans": ("performans", "performans notu", "performans sonucu")}
TALEP_SUTUNLARI = {"sicil": ("sicil no", "sicil"), "talep": ("talep edilen egitim", "talep", "egitim talebi")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def seviye(x) -> float | None:
    if x in (None, ""):
        return None
    try:
        v = float(str(x).replace(",", "."))
    except ValueError:
        return None
    return v if 0 <= v <= 5 else None


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Gereklilik:
    pozisyon: str
    yetkinlik: str
    gerekli: float
    kritiklik: int
    egitim: str


@dataclass
class Ihtiyac:
    sicil: str
    ad: str
    departman: str
    pozisyon: str
    g: Gereklilik
    yonetici: float | None
    oz: float | None
    performans: str = ""
    talep: bool = False
    notlar: list = field(default_factory=list)

    @property
    def mevcut(self) -> float | None:
        return self.yonetici if self.yonetici is not None else self.oz

    @property
    def acik(self) -> float:
        return max(self.g.gerekli - self.mevcut, 0) if self.mevcut is not None else 0

    @property
    def dusuk_perf(self) -> bool:
        return katla(self.performans)[:1] in DUSUK_PERF if self.performans else False

    @property
    def puan(self) -> float:
        if self.acik <= 0:
            return 0
        return self.acik * self.g.kritiklik + self.dusuk_perf + self.talep

    @property
    def oncelik(self) -> int | None:
        p = self.puan
        return None if p <= 0 else 1 if p >= 6 else 2 if p >= 3 else 3


def oku_matris(yol: Path) -> tuple[list[Gereklilik], list[dict]]:
    sonuc, uy = [], []
    for r in kayitlar(yol, MATRIS_SUTUNLARI, ("pozisyon", "yetkinlik", "gerekli")):
        g = seviye(r.get("gerekli"))
        if not r.get("pozisyon") or not r.get("yetkinlik") or g is None:
            uy.append({"onem": "Orta", "tur": "Matris satırı okunamadı", "kim": "", "aciklama": f"{yol.name} satır {r['_satir']}"})
            continue
        sonuc.append(Gereklilik(metin(r["pozisyon"]), metin(r["yetkinlik"]), g, KRITIKLIK.get(katla(r.get("kritiklik")), 2),
                                metin(r.get("egitim")) or metin(r["yetkinlik"])))
    return sonuc, uy


def oku_tablo(yol: Path | None, sozluk, zorunlu) -> list[dict]:
    return kayitlar(yol, sozluk, zorunlu) if yol else []


# ----------------------------------------------------------------------------
# Analiz
# ----------------------------------------------------------------------------

def analiz_et(matris: list[Gereklilik], degerlendirmeler: list[dict], performans: dict[str, str], talepler: list[dict],
              min_grup: int = 5, fark_esik: float = 2) -> dict:
    uyarilar = []

    def uyar(onem, tur, kim, aciklama):
        uyarilar.append({"onem": onem, "tur": tur, "kim": kim, "aciklama": aciklama})

    gerek = defaultdict(dict)
    for g in matris:
        gerek[katla(g.pozisyon)][katla(g.yetkinlik)] = g
    kisiler: dict[str, dict] = {}
    puanlar: dict[tuple[str, str], dict] = {}
    for r in degerlendirmeler:
        s = metin(r.get("sicil"))
        if not s:
            continue
        k = kisiler.setdefault(s, {"ad": metin(r.get("ad")), "departman": metin(r.get("departman")) or "—", "pozisyon": metin(r.get("pozisyon"))})
        puanlar[(s, katla(r.get("yetkinlik")))] = {"ad": metin(r.get("yetkinlik")), "yonetici": seviye(r.get("yonetici")), "oz": seviye(r.get("oz"))}
        if katla(k["pozisyon"]) not in gerek:
            continue
        if katla(r.get("yetkinlik")) not in gerek[katla(k["pozisyon"])]:
            uyar("Bilgi", "Matriste olmayan yetkinlik", s, f"'{metin(r.get('yetkinlik'))}' {k['pozisyon']} matrisinde yok; değerlendirildi ama analize alınmadı")
    ihtiyaclar: list[Ihtiyac] = []
    for s, k in kisiler.items():
        poz = katla(k["pozisyon"])
        if poz not in gerek:
            uyar("Orta", "Pozisyon matriste yok", s, f"'{k['pozisyon']}' için yetkinlik matrisi tanımlı değil")
            continue
        for y, g in gerek[poz].items():
            p = puanlar.get((s, y))
            if not p or (p["yonetici"] is None and p["oz"] is None):
                uyar("Orta", "Değerlendirilmedi", s, f"'{g.yetkinlik}' için yönetici veya öz değerlendirme puanı yok")
                continue
            i = Ihtiyac(s, k["ad"], k["departman"], k["pozisyon"], g, p["yonetici"], p["oz"], performans.get(s, ""))
            if p["yonetici"] is None:
                i.notlar.append("Yalnız öz değerlendirme")
            ihtiyaclar.append(i)
    if performans:
        for s in kisiler:
            if s not in performans:
                uyar("Bilgi", "Performans sonucu yok", s, "Performans dosyasında kayıt yok")
    # Talepler
    talep_sonuc = []
    for r in talepler:
        s, t = metin(r.get("sicil")), metin(r.get("talep"))
        if not s or not t:
            continue
        kt = katla(t)
        eslesen = [i for i in ihtiyaclar if i.sicil == s and kt and (kt in katla(i.g.egitim) or katla(i.g.egitim) in kt
                                                                      or kt in katla(i.g.yetkinlik) or katla(i.g.yetkinlik) in kt)]
        for i in eslesen:
            i.talep = True
        acikli = [i for i in eslesen if i.acik > 0]
        durum = ("Talep ve ihtiyaç örtüşüyor" if acikli else "Talep var, ölçülen açık yok (gelişim talebi)" if eslesen
                 else "Matristeki eğitimlerle eşleşmedi" if s in kisiler else "Çalışan değerlendirmelerde yok")
        talep_sonuc.append({"sicil": s, "ad": kisiler.get(s, {}).get("ad", ""), "talep": t, "durum": durum,
                            "yetkinlik": ", ".join(i.g.yetkinlik for i in eslesen)})
    # Algı farkları
    farklar = []
    for i in ihtiyaclar:
        if i.yonetici is not None and i.oz is not None and abs(i.oz - i.yonetici) >= fark_esik:
            tur = "Kör nokta (öz değerlendirme yüksek)" if i.oz > i.yonetici else "Fark edilmemiş güç (öz değerlendirme düşük)"
            farklar.append({"i": i, "tur": tur, "fark": i.oz - i.yonetici})
            i.notlar.append(tur.split(" (")[0])
    # Departman × yetkinlik önceliği
    grup = defaultdict(list)
    for i in ihtiyaclar:
        grup[(i.departman, i.g.yetkinlik)].append(i)
    oncelik = []
    for (d, y), lst in grup.items():
        acikli = [i for i in lst if i.acik > 0]
        oncelik.append({"departman": d, "yetkinlik": y, "kritiklik": lst[0].g.kritiklik, "egitim": lst[0].g.egitim, "degerlendirilen": len(lst),
                        "acikli": len(acikli), "oran": len(acikli) / len(lst), "ort_acik": statistics.fmean(i.acik for i in lst),
                        "puan": sum(i.puan for i in lst), "o1": sum(1 for i in lst if i.oncelik == 1)})
    oncelik.sort(key=lambda x: (-x["puan"], -x["kritiklik"], x["departman"], x["yetkinlik"]))
    for n, x in enumerate(oncelik, 1):
        x["sira"] = n
    # Eğitim planı
    egitim = defaultdict(list)
    for i in ihtiyaclar:
        if i.acik > 0:
            egitim[i.g.egitim].append(i)
    plan = []
    for ad, lst in egitim.items():
        n = len({i.sicil for i in lst})
        plan.append({"egitim": ad, "kisi": n, "departmanlar": ", ".join(sorted({i.departman for i in lst})), "puan": sum(i.puan for i in lst),
                     "o1": sum(1 for i in lst if i.oncelik == 1), "yontem": "Grup eğitimi (sınıf içi / sanal sınıf)" if n >= min_grup
                     else "Bireysel: e-öğrenme, mentorluk veya iş başında eğitim", "katilimcilar": sorted(lst, key=lambda i: (i.oncelik, i.sicil))})
    plan.sort(key=lambda x: (-x["puan"], x["egitim"]))
    # Departman özeti
    dep = defaultdict(list)
    for i in ihtiyaclar:
        dep[i.departman].append(i)
    ozet = [{"departman": d, "kisi": len({i.sicil for i in lst}), "degerlendirme": len(lst), "karsilanan": sum(1 for i in lst if i.acik <= 0),
             "oran": sum(1 for i in lst if i.acik <= 0) / len(lst), "o1": sum(1 for i in lst if i.oncelik == 1),
             "o2": sum(1 for i in lst if i.oncelik == 2), "o3": sum(1 for i in lst if i.oncelik == 3)} for d, lst in sorted(dep.items())]
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    uyarilar.sort(key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))
    return {"ihtiyaclar": ihtiyaclar, "oncelik": oncelik, "plan": plan, "ozet": ozet, "farklar": farklar, "talepler": talep_sonuc, "uyarilar": uyarilar,
            "min_grup": min_grup}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
ONCELIK_RENK = {1: "FDE2E1", 2: "FFF4CE", 3: "E8F0FE"}
KONTROL = PatternFill("solid", fgColor="FFF4CE")
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def rapor_yaz(cikti: Path, s: dict) -> None:
    wb = Workbook()
    ol = wb.active
    ol.title = "Öncelik Listesi"
    _baslik(ol, ["Sıra", "Departman", "Yetkinlik", "Kritiklik", "Önerilen Eğitim", "Değerlendirilen", "Açığı Olan", "Açık Oranı", "Ort. Açık",
                 "İhtiyaç Puanı", "Öncelik 1 Kişi", "Karar"], (6, 14, 24, 9, 32, 12, 10, 10, 9, 10, 10, 20))
    for x in s["oncelik"]:
        ol.append([x["sira"], x["departman"], x["yetkinlik"], KRITIK_AD[x["kritiklik"]], x["egitim"], x["degerlendirilen"], x["acikli"], x["oran"],
                   round(x["ort_acik"], 2), x["puan"], x["o1"], ""])
        ol.cell(ol.max_row, 8).number_format = "0%"
        if x["o1"]:
            ol.cell(ol.max_row, 11).fill = PatternFill("solid", fgColor="FDE2E1")
        ol.cell(ol.max_row, 12).fill = KONTROL

    ep = wb.create_sheet("Eğitim Planı")
    _baslik(ep, ["Önerilen Eğitim", "Katılımcı", "Departmanlar", "Öncelik 1", "Toplam Puan", "Önerilen Yöntem", "Katılımcılar (sicil · öncelik)",
                 "Planlanan Tarih", "Bütçe"], (32, 10, 22, 9, 10, 36, 50, 14, 12))
    for x in s["plan"]:
        ep.append([x["egitim"], x["kisi"], x["departmanlar"], x["o1"], x["puan"], x["yontem"],
                   ", ".join(f"{i.sicil} · Ö{i.oncelik}" for i in x["katilimcilar"]), "", ""])
        ep.cell(ep.max_row, 7).alignment = UST
        ep.cell(ep.max_row, 8).fill = ep.cell(ep.max_row, 9).fill = KONTROL
    ep.append([])
    ep.append([f"Katılımcı sayısı {s['min_grup']} ve üstü olan eğitimler için grup eğitimi önerilir (--min-grup)."])

    bi = wb.create_sheet("Bireysel İhtiyaçlar")
    _baslik(bi, ["Sicil", "Ad Soyad", "Departman", "Pozisyon", "Yetkinlik", "Kritiklik", "Gerekli", "Yönetici", "Öz", "Açık", "Performans", "Talep",
                 "Puan", "Öncelik", "Önerilen Eğitim", "Not", "Onay"], (8, 16, 12, 16, 24, 9, 8, 9, 6, 6, 10, 7, 7, 8, 30, 26, 10))
    for i in sorted(s["ihtiyaclar"], key=lambda i: (i.oncelik or 9, -i.puan, i.sicil)):
        if i.acik <= 0 and not i.notlar:
            continue
        bi.append([i.sicil, i.ad, i.departman, i.pozisyon, i.g.yetkinlik, KRITIK_AD[i.g.kritiklik], i.g.gerekli, i.yonetici, i.oz, i.acik,
                   i.performans, "Evet" if i.talep else "", i.puan, i.oncelik, i.g.egitim, "; ".join(i.notlar), ""])
        if i.oncelik:
            bi.cell(bi.max_row, 14).fill = PatternFill("solid", fgColor=ONCELIK_RENK[i.oncelik])
        bi.cell(bi.max_row, 17).fill = KONTROL
    bi.auto_filter.ref = f"A1:Q{bi.max_row}"

    dz = wb.create_sheet("Departman Özeti")
    _baslik(dz, ["Departman", "Çalışan", "Değerlendirme", "Karşılanan", "Yetkinlik Karşılama Oranı", "Öncelik 1", "Öncelik 2", "Öncelik 3"],
            (16, 9, 12, 11, 14, 9, 9, 9))
    for x in s["ozet"]:
        dz.append([x["departman"], x["kisi"], x["degerlendirme"], x["karsilanan"], x["oran"], x["o1"], x["o2"], x["o3"]])
        dz.cell(dz.max_row, 5).number_format = "0%"
    dz.append([])
    for t in ("Puan = açık × kritiklik (Yüksek 3, Orta 2, Düşük 1) + 1 (performans C/D veya 1–2) + 1 (çalışan talebi). "
              "Öncelik: puan ≥ 6 → 1, 3–5 → 2, 1–2 → 3.",
              "Mevcut seviye yönetici puanıdır; yönetici puanı yoksa öz değerlendirme kullanılır ve 'Not' sütununda belirtilir."):
        dz.append([t])

    af = wb.create_sheet("Algı Farkları")
    _baslik(af, ["Sicil", "Ad Soyad", "Departman", "Yetkinlik", "Yönetici", "Öz", "Fark", "Tür", "Görüşme Notu"], (8, 16, 12, 24, 9, 6, 6, 40, 30))
    for x in s["farklar"]:
        i = x["i"]
        af.append([i.sicil, i.ad, i.departman, i.g.yetkinlik, i.yonetici, i.oz, x["fark"], x["tur"], ""])
        af.cell(af.max_row, 9).fill = KONTROL

    tl_ = wb.create_sheet("Talepler")
    _baslik(tl_, ["Sicil", "Ad Soyad", "Talep Edilen Eğitim", "Eşleşen Yetkinlik", "Durum"], (8, 16, 32, 26, 44))
    for x in s["talepler"]:
        tl_.append([x["sicil"], x["ad"], x["talep"], x["yetkinlik"], x["durum"]])

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Kim", "Açıklama"], (9, 28, 10, 90))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["kim"], u["aciklama"]])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(matris_yolu: Path, deger_yolu: Path, cikti: Path, perf_yolu: Path | None = None, talep_yolu: Path | None = None,
             min_grup: int = 5) -> dict:
    matris, uy = oku_matris(matris_yolu)
    if not matris:
        raise ValueError(f"{matris_yolu.name}: yetkinlik matrisi boş")
    perf = {metin(r.get("sicil")): metin(r.get("performans")) for r in oku_tablo(perf_yolu, PERF_SUTUNLARI, ("sicil", "performans")) if r.get("sicil")}
    s = analiz_et(matris, oku_tablo(deger_yolu, DEGER_SUTUNLARI, ("sicil", "yetkinlik")), perf,
                  oku_tablo(talep_yolu, TALEP_SUTUNLARI, ("sicil", "talep")), min_grup)
    s["uyarilar"] = uy + s["uyarilar"]
    rapor_yaz(cikti, s)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Yetkinlik değerlendirmeleri, performans ve taleplerden önceliklendirilmiş eğitim ihtiyacı çıkarır.")
    p.add_argument("--matris", type=Path, default=ORNEK / "yetkinlik_matrisi.csv", help="Pozisyon, Yetkinlik, Gerekli Seviye, Kritiklik, Önerilen Eğitim")
    p.add_argument("--degerlendirmeler", type=Path, default=ORNEK / "degerlendirmeler.csv",
                   help="Sicil, Ad, Departman, Pozisyon, Yetkinlik, Yönetici Puanı, Öz Değerlendirme (1–5)")
    p.add_argument("--performans", type=Path, help="İsteğe bağlı: Sicil, Performans (A–D veya 1–5)")
    p.add_argument("--talepler", type=Path, help="İsteğe bağlı: Sicil, Talep Edilen Eğitim")
    p.add_argument("--min-grup", type=int, default=5, help="Grup eğitimi için en az katılımcı (varsayılan 5)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "egitim_ihtiyac_analizi.xlsx")
    a = p.parse_args(argv)
    ornek = a.degerlendirmeler == ORNEK / "degerlendirmeler.csv"
    perf = a.performans or (ORNEK / "performans.csv" if ornek else None)
    talep = a.talepler or (ORNEK / "talepler.csv" if ornek else None)
    for y in (a.matris, a.degerlendirmeler, perf, talep):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        s = calistir(a.matris, a.degerlendirmeler, a.cikti, perf, talep, a.min_grup)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    acik = [i for i in s["ihtiyaclar"] if i.acik > 0]
    print(f"[OK] {len({i.sicil for i in s['ihtiyaclar']})} çalışan · {len(s['ihtiyaclar'])} yetkinlik değerlendirmesi · {len(acik)} açık "
          f"(öncelik 1: {sum(1 for i in acik if i.oncelik == 1)})")
    for x in s["oncelik"][:5]:
        print(f"     {x['sira']}. {x['departman']} · {x['yetkinlik']}: {x['acikli']}/{x['degerlendirilen']} kişide açık, puan {x['puan']:g}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
