import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "h.xlsx"
        cls.s = main.calistir(ORNEK / "kkn_tanimlari.csv", ORNEK / "izleme_kayitlari.csv", cls.cikti)
        cls.u = {}
        for u in cls.s["uyarilar"]:
            cls.u.setdefault(u["tur"], []).append(u)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_plan(self):
        p = self.s["plan"]
        self.assertEqual((p["KKN-1"].k_alt, p["KKN-1"].o_alt, p["KKN-3"].k_ust, p["KKN-3"].o_ust), (D(72), D(73), D(4), D("3.5")))
        self.assertTrue(p["KKN-4"].kategorik)

    def test_kritik_sapmalar(self):
        s = [(x["kkn"].kod, x["kayit"].ham, x["df_var"]) for x in self.s["sapmalar"]]
        self.assertEqual(s, [("KKN-1", "71,4", True), ("KKN-1", "71,8", False), ("KKN-3", "4,3", True), ("KKN-4", "Kaldı", True)])
        self.assertEqual([(x["kkn"].kod, x["kayit"], x["kritik"], x["df_yok"]) for x in self.s["ozet"]],
                         [("KKN-1", 36, 2, 1), ("KKN-2", 20, 0, 0), ("KKN-3", 20, 1, 0), ("KKN-4", 10, 1, 0)])

    def test_uyarilar(self):
        self.assertIn("10:30 – 13:00 arası 150 dk", self.u["İzleme boşluğu"][0]["aciklama"])
        self.assertEqual(len(self.u["Operasyonel limit dışı"]), 3)
        self.assertIn("2,7 → 2,9 → 3,1 → 3,3 → 3,6", self.u["Limite yaklaşan eğilim"][0]["aciklama"])
        self.assertIn("art arda 8 kez 74,0", self.u["Tekrar eden aynı değer"][0]["aciklama"])
        self.assertEqual(len(self.u["Doğrulama yok"]), 1)
        self.assertNotIn("Kayıt yok", self.u)

    def test_parametreler(self):
        s = main.calistir(ORNEK / "kkn_tanimlari.csv", ORNEK / "izleme_kayitlari.csv", self.cikti, egilim_n=8, tekrar_n=9)
        turler = {u["tur"] for u in s["uyarilar"]}
        self.assertNotIn("Limite yaklaşan eğilim", turler)
        self.assertNotIn("Tekrar eden aynı değer", turler)

    def test_excel(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Özet", "Sapmalar", "Uyarılar", "Kayıtlar", "Grafikler", "HACCP Planı"])
        sp = wb["Sapmalar"]
        self.assertEqual(sp["F3"].value, "KAYIT YOK")
        self.assertEqual(sp["H1"].value, "Ürün Durumu / QA Değerlendirmesi")
        self.assertEqual(len(wb["Grafikler"]._charts), 3)


class KuralTesti(unittest.TestCase):
    def test_eksik_kkn_ve_planda_olmayan(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t / "p.csv").write_text("KKN Kodu;Parametre;Kritik Üst\nA;pH;4,6\nB;Sıcaklık;4\n", encoding="utf-8")
            (t / "k.csv").write_text("Tarih;Saat;KKN Kodu;Değer\n01.10.2026;08:00;A;4,4\n01.10.2026;09:00;A;yok\n01.10.2026;09:00;Z;1\n", encoding="utf-8")
            s = main.calistir(t / "p.csv", t / "k.csv", t / "o.xlsx")
            turler = [(u["tur"], u["kkn"]) for u in s["uyarilar"]]
            for beklenen in (("Kayıt yok", "B"), ("Planda olmayan KKN", "Z"), ("Okunamayan değer", "A")):
                self.assertIn(beklenen, turler)
            self.assertEqual(s["sapmalar"], [])

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--plan", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
