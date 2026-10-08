"""
HACCP Kritik Kontrol Noktası İzleme Analizi — Workers / Workless kod bloğu
Gıda Üretimi › Kalite Güvence ve Gıda Güvenliği › Kalite Güvence Uzmanı

KKN izleme kayıtlarını (sıcaklık, süre, pH, metal dedektör…) HACCP planındaki kritik ve operasyonel limitlerle
karşılaştırır:
  - Kritik limit sapması (Kritik): düzeltici faaliyet kaydı yoksa ayrıca işaretlenir; etkilenen parti/lot listelenir.
  - Operasyonel limit (hedef) aşımı: kritik limite ulaşmadan müdahale gerektiren uyarı.
  - Limite yaklaşan eğilim: art arda N ölçüm (varsayılan 5) sürekli limit yönünde.
  - İzleme boşluğu: aynı gün içinde iki kayıt arası, izleme sıklığının 1,5 katını aşıyor.
  - Kayıt bütünlüğü: art arda çok sayıda birebir aynı değer (varsayılan 8), ölçen / doğrulayan boş.
  - Kategorik KKN'ler (metal dedektör test kartı): "Kaldı / Red / Uygunsuz" sonuçları kritik sapma sayılır.
Rapor: KKN bazında uyum özeti, sapmalar ('QA Değerlendirmesi'), uyarılar, işaretli kayıtlar, sayısal KKN grafikleri.
İnternete bağlanmaz.

Kullanım:
    python main.py                                                     # örnek süt işletmesi, 2 gün kayıt
    python main.py --plan kkn_tanimlari.xlsx --kayitlar izleme.xlsx --egilim 5 --tekrar 8
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
OLUMSUZ_SONUC = ("kaldi", "red", "uygunsuz", "basarisiz", "fail", "tespit edilemedi", "hayir")
OLUMLU_SONUC = ("gecti", "uygun", "basarili", "pass", "ok", "evet")

PLAN_SUTUNLARI = {"kod": ("kkn kodu", "kkn", "kod"), "adim": ("adim", "proses adimi"), "parametre": ("parametre", "izlenen parametre"),
                  "birim": ("birim",), "k_alt": ("kritik alt", "kritik limit alt", "alt kritik limit"),
                  "k_ust": ("kritik ust", "kritik limit ust", "ust kritik limit"), "o_alt": ("operasyonel alt", "hedef alt", "uyari alt"),
                  "o_ust": ("operasyonel ust", "hedef ust", "uyari ust"), "siklik": ("izleme sikligi dk", "izleme sikligi", "siklik dk", "siklik"),
                  "df": ("tanimli duzeltici faaliyet", "duzeltici faaliyet")}
KAYIT_SUTUNLARI = {"tarih": ("tarih",), "saat": ("saat",), "kod": ("kkn kodu", "kkn", "kod"), "deger": ("deger", "olcum", "sonuc"),
                   "lot": ("parti lot", "parti", "lot", "parti no"), "olcen": ("olcen", "kaydeden", "operator"),
                   "df": ("duzeltici faaliyet", "yapilan duzeltici faaliyet", "df"), "dogrulayan": ("dogrulayan", "kontrol eden", "onaylayan")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def sayi(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace(",", ".")
    if not re.fullmatch(r"-?\d+(\.\d+)?", s):
        return None
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def goster(x) -> str:
    return "—" if x is None else f"{x:g}".replace(".", ",")


def zaman(t, s) -> datetime | None:
    if isinstance(t, datetime) and not s:
        return t
    d = t.date() if isinstance(t, datetime) else t if isinstance(t, date) else None
    if d is None:
        for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
            try:
                d = datetime.strptime(str(t or "").strip()[:10], f).date()
                break
            except ValueError:
                pass
    if d is None:
        return None
    if isinstance(s, time):
        return datetime.combine(d, s)
    m = re.search(r"(\d{1,2})[:.](\d{2})", str(s or "") or str(t or "")[10:])
    return datetime.combine(d, time(int(m.group(1)), int(m.group(2))) if m else time(0, 0))


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class KKN:
    kod: str
    adim: str
    parametre: str
    birim: str
    k_alt: Decimal | None
    k_ust: Decimal | None
    o_alt: Decimal | None
    o_ust: Decimal | None
    siklik: int | None
    df: str

    @property
    def kategorik(self) -> bool:
        return self.k_alt is None and self.k_ust is None

    @property
    def ad(self) -> str:
        return f"{self.kod} {self.adim} — {self.parametre}" + (f" ({self.birim})" if self.birim else "")

    def limit_metni(self) -> str:
        if self.kategorik:
            return "Geçti olmalı"
        p = []
        if self.k_alt is not None:
            p.append(f"≥ {goster(self.k_alt)}")
        if self.k_ust is not None:
            p.append(f"≤ {goster(self.k_ust)}")
        return " ve ".join(p) + (f" {self.birim}" if self.birim else "")


@dataclass
class Kayit:
    satir: int
    an: datetime
    kod: str
    ham: str
    deger: Decimal | None
    lot: str
    olcen: str
    df: str
    dogrulayan: str
    durum: str = "Uygun"
    isaretler: list = field(default_factory=list)


def plan_oku(yol: Path) -> dict[str, KKN]:
    return {str(r["kod"]).strip(): KKN(str(r["kod"]).strip(), str(r.get("adim") or "").strip(), str(r.get("parametre") or "").strip(),
                                       str(r.get("birim") or "").strip(), sayi(r.get("k_alt")), sayi(r.get("k_ust")), sayi(r.get("o_alt")),
                                       sayi(r.get("o_ust")), int(sayi(r.get("siklik"))) if sayi(r.get("siklik")) else None, str(r.get("df") or "").strip())
            for r in kayitlar(yol, PLAN_SUTUNLARI, ("kod",)) if r.get("kod")}


def kayitlari_oku(yol: Path) -> tuple[list[Kayit], list[str]]:
    sonuc, hatalar = [], []
    for r in kayitlar(yol, KAYIT_SUTUNLARI, ("tarih", "kod", "deger")):
        if not r.get("kod"):
            continue
        an = zaman(r.get("tarih"), r.get("saat"))
        if an is None:
            hatalar.append(f"Satır {r['_satir']}: tarih okunamadı")
            continue
        sonuc.append(Kayit(r["_satir"], an, str(r["kod"]).strip(), str(r.get("deger") if r.get("deger") is not None else "").strip(), sayi(r.get("deger")),
                           str(r.get("lot") or "").strip(), str(r.get("olcen") or "").strip(), str(r.get("df") or "").strip(),
                           str(r.get("dogrulayan") or "").strip()))
    return sorted(sonuc, key=lambda k: (k.kod, k.an)), hatalar


# ----------------------------------------------------------------------------
# Analiz
# ----------------------------------------------------------------------------

def analiz_et(plan: dict[str, KKN], kayit_listesi: list[Kayit], egilim_n: int = 5, tekrar_n: int = 8) -> dict:
    sapmalar, uyarilar = [], []

    def uyari(onem, tur, kkn, aciklama, ilgili=()):
        uyarilar.append({"onem": onem, "tur": tur, "kkn": kkn, "aciklama": aciklama, "kayitlar": [k.satir for k in ilgili]})
        for k in ilgili:
            k.isaretler.append(tur)

    gruplar = defaultdict(list)
    for k in kayit_listesi:
        gruplar[k.kod].append(k)
    for kod in sorted(set(gruplar) - set(plan)):
        uyari("Orta", "Planda olmayan KKN", kod, f"{kod}: {len(gruplar[kod])} kayıt var ama HACCP planında tanımı yok", gruplar[kod][:1])
    for kod, kkn in plan.items():
        lst = gruplar.get(kod, [])
        if not lst:
            uyari("Yüksek", "Kayıt yok", kod, f"{kkn.ad}: dönemde hiç izleme kaydı yok")
            continue
        for k in lst:
            if kkn.kategorik:
                kk = katla(k.ham)
                if any(kk.startswith(x) for x in OLUMSUZ_SONUC):
                    k.durum = "Kritik sapma"
                elif not any(kk.startswith(x) for x in OLUMLU_SONUC):
                    k.durum = "Okunamadı"
            elif k.deger is None:
                k.durum = "Okunamadı"
            elif (kkn.k_alt is not None and k.deger < kkn.k_alt) or (kkn.k_ust is not None and k.deger > kkn.k_ust):
                k.durum = "Kritik sapma"
            elif (kkn.o_alt is not None and k.deger < kkn.o_alt) or (kkn.o_ust is not None and k.deger > kkn.o_ust):
                k.durum = "Operasyonel limit dışı"
            if k.durum == "Kritik sapma":
                sapmalar.append({"kkn": kkn, "kayit": k, "df_var": bool(k.df)})
                k.isaretler.append("Kritik sapma" + ("" if k.df else " — DF kaydı yok"))
            elif k.durum == "Operasyonel limit dışı":
                uyari("Orta", "Operasyonel limit dışı", kod, f"{k.an:%d.%m.%Y %H:%M} {goster(k.deger)} {kkn.birim} (operasyonel limit "
                      f"{'≥ ' + goster(kkn.o_alt) if kkn.o_alt is not None else '≤ ' + goster(kkn.o_ust)}; kritik limit {kkn.limit_metni()})", [k])
            elif k.durum == "Okunamadı":
                uyari("Orta", "Okunamayan değer", kod, f"{k.an:%d.%m.%Y %H:%M}: '{k.ham}' değeri yorumlanamadı", [k])
            if not k.olcen:
                uyari("Bilgi", "Ölçen boş", kod, f"{k.an:%d.%m.%Y %H:%M}: ölçen / kaydeden yazılmamış", [k])
            if not k.dogrulayan:
                uyari("Bilgi", "Doğrulama yok", kod, f"{k.an:%d.%m.%Y %H:%M}: doğrulayan imzası yok", [k])
        if kkn.siklik:
            for a, b in zip(lst, lst[1:]):
                fark = (b.an - a.an).total_seconds() / 60
                if a.an.date() == b.an.date() and fark > kkn.siklik * 1.5:
                    uyari("Yüksek", "İzleme boşluğu", kod, f"{a.an:%d.%m.%Y} {a.an:%H:%M} – {b.an:%H:%M} arası {int(fark)} dk kayıt yok "
                          f"(izleme sıklığı {kkn.siklik} dk)", [a, b])
        if not kkn.kategorik:
            yon = 1 if kkn.k_ust is not None else -1          # +1: üst limite doğru artış tehlikeli
            seri = []
            bildirilen = False
            for k in lst:
                if seri and k.an.date() == seri[-1].an.date() and k.deger is not None and seri[-1].deger is not None and (k.deger - seri[-1].deger) * yon > 0:
                    seri.append(k)
                else:
                    seri, bildirilen = [k], False
                if len(seri) >= egilim_n and not bildirilen:
                    uyari("Orta", "Limite yaklaşan eğilim", kod, f"{seri[0].an:%d.%m.%Y %H:%M} itibarıyla art arda {len(seri)} ölçüm "
                          f"{'artıyor' if yon > 0 else 'azalıyor'}: " + " → ".join(goster(x.deger) for x in seri), list(seri))
                    bildirilen = True
            ayni = []
            for k in lst + [None]:
                if k is not None and ayni and k.ham == ayni[-1].ham and k.an.date() == ayni[-1].an.date():
                    ayni.append(k)
                    continue
                if len(ayni) >= tekrar_n:
                    uyari("Orta", "Tekrar eden aynı değer", kod, f"{ayni[0].an:%d.%m.%Y %H:%M} – {ayni[-1].an:%H:%M} arası art arda {len(ayni)} kez "
                          f"{ayni[0].ham}: kayıtların gerçek ölçümle tutulduğunu doğrulayın", ayni)
                ayni = [k] if k is not None else []
    ozet = []
    for kod, kkn in plan.items():
        lst = gruplar.get(kod, [])
        kritik = sum(k.durum == "Kritik sapma" for k in lst)
        ozet.append({"kkn": kkn, "kayit": len(lst), "kritik": kritik, "df_yok": sum(1 for s in sapmalar if s["kkn"] is kkn and not s["df_var"]),
                     "operasyonel": sum(k.durum == "Operasyonel limit dışı" for k in lst),
                     "uyum": (len(lst) - kritik) / len(lst) if lst else None,
                     "min": min((k.deger for k in lst if k.deger is not None), default=None), "max": max((k.deger for k in lst if k.deger is not None), default=None)})
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    uyarilar.sort(key=lambda u: sira[u["onem"]])
    return {"sapmalar": sapmalar, "uyarilar": uyarilar, "ozet": ozet, "gruplar": gruplar}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE", "Kritik sapma": "FDE2E1", "Operasyonel limit dışı": "FFF4CE", "Okunamadı": "EEEEEE"}
KONTROL = PatternFill("solid", fgColor="FFF4CE")
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def rapor_yaz(cikti: Path, plan: dict[str, KKN], kayit_listesi: list[Kayit], s: dict, hatalar: list[str]) -> None:
    wb = Workbook()
    o = wb.active
    o.title = "Özet"
    _baslik(o, ["KKN", "Kritik Limit", "Kayıt", "Kritik Sapma", "DF Kaydı Yok", "Operasyonel Limit Dışı", "Uyum", "En Düşük", "En Yüksek"],
            (44, 18, 8, 10, 10, 12, 9, 10, 10))
    for x in s["ozet"]:
        o.append([x["kkn"].ad, x["kkn"].limit_metni(), x["kayit"], x["kritik"], x["df_yok"], x["operasyonel"], x["uyum"],
                  x["min"] and float(x["min"]), x["max"] and float(x["max"])])
        o.cell(o.max_row, 7).number_format = "0.0%"
        if x["kritik"]:
            o.cell(o.max_row, 4).fill = PatternFill("solid", fgColor=RENK["Yüksek"])
        if x["df_yok"]:
            o.cell(o.max_row, 5).fill = PatternFill("solid", fgColor=RENK["Yüksek"])
    o.append([])
    o.append(["Kayıt dönemi", f"{min(k.an for k in kayit_listesi):%d.%m.%Y %H:%M} – {max(k.an for k in kayit_listesi):%d.%m.%Y %H:%M}" if kayit_listesi else "—"])
    for h in hatalar:
        o.append(["Okunamayan satır", h])
    o.append(["Not", "Kritik sapmalarda etkilenen ürünün durumu, düzeltici faaliyetin yeterliliği ve kök neden QA tarafından değerlendirilmelidir."])

    sp = wb.create_sheet("Sapmalar")
    _baslik(sp, ["Tarih / Saat", "KKN", "Değer", "Kritik Limit", "Parti / Lot", "Kaydedilen Düzeltici Faaliyet", "Tanımlı Düzeltici Faaliyet",
                 "Ürün Durumu / QA Değerlendirmesi"], (16, 36, 9, 16, 11, 50, 44, 34))
    for x in s["sapmalar"]:
        k, kkn = x["kayit"], x["kkn"]
        sp.append([k.an, kkn.ad, k.ham, kkn.limit_metni(), k.lot, k.df or "KAYIT YOK", kkn.df, ""])
        sp.cell(sp.max_row, 1).number_format = "DD.MM.YYYY HH:MM"
        sp.cell(sp.max_row, 3).fill = PatternFill("solid", fgColor=RENK["Yüksek"])
        if not k.df:
            sp.cell(sp.max_row, 6).fill = PatternFill("solid", fgColor=RENK["Yüksek"])
            sp.cell(sp.max_row, 6).font = Font(bold=True)
        sp.cell(sp.max_row, 8).fill = KONTROL
        for c in sp[sp.max_row]:
            c.alignment = UST

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "KKN", "Açıklama", "Kayıt Satırları", "İnceleme"], (9, 26, 9, 90, 16, 24))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["kkn"], u["aciklama"], ", ".join(map(str, u["kayitlar"][:10])) + (" …" if len(u["kayitlar"]) > 10 else ""), ""])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 6).fill = KONTROL
        uy.cell(uy.max_row, 4).alignment = UST

    ky = wb.create_sheet("Kayıtlar")
    _baslik(ky, ["Satır", "Tarih / Saat", "KKN", "Değer", "Parti / Lot", "Ölçen", "Düzeltici Faaliyet", "Doğrulayan", "Durum", "İşaretler"],
            (6, 16, 8, 9, 10, 18, 40, 16, 20, 40))
    for k in sorted(kayit_listesi, key=lambda k: (k.an, k.kod)):
        ky.append([k.satir, k.an, k.kod, float(k.deger) if k.deger is not None else k.ham, k.lot, k.olcen, k.df, k.dogrulayan, k.durum,
                   "; ".join(dict.fromkeys(k.isaretler))])
        ky.cell(ky.max_row, 2).number_format = "DD.MM.YYYY HH:MM"
        if k.durum in RENK:
            ky.cell(ky.max_row, 9).fill = PatternFill("solid", fgColor=RENK[k.durum])
    ky.auto_filter.ref = f"A1:J{ky.max_row}"

    gr = wb.create_sheet("Grafikler")
    satir = 1
    for kod, kkn in plan.items():
        lst = [k for k in s["gruplar"].get(kod, []) if k.deger is not None]
        if kkn.kategorik or not lst:
            continue
        gr.cell(satir, 1, kkn.ad).font = Font(bold=True)
        gr.cell(satir + 1, 1, "Zaman")
        gr.cell(satir + 1, 2, "Değer")
        limitler = [(a, v) for a, v in (("Kritik alt", kkn.k_alt), ("Kritik üst", kkn.k_ust), ("Operasyonel alt", kkn.o_alt), ("Operasyonel üst", kkn.o_ust)) if v is not None]
        for j, (a, _) in enumerate(limitler, 3):
            gr.cell(satir + 1, j, a)
        for i, k in enumerate(lst, satir + 2):
            gr.cell(i, 1, k.an.strftime("%d.%m %H:%M"))
            gr.cell(i, 2, float(k.deger))
            for j, (_, v) in enumerate(limitler, 3):
                gr.cell(i, j, float(v))
        ch = LineChart()
        ch.title, ch.height, ch.width = kkn.ad, 7, 22
        ch.add_data(Reference(gr, min_col=2, max_col=2 + len(limitler), min_row=satir + 1, max_row=satir + 1 + len(lst)), titles_from_data=True)
        ch.set_categories(Reference(gr, min_col=1, min_row=satir + 2, max_row=satir + 1 + len(lst)))
        gr.add_chart(ch, f"H{satir}")
        satir += max(len(lst) + 4, 17)

    pl = wb.create_sheet("HACCP Planı")
    _baslik(pl, ["KKN", "Adım", "Parametre", "Birim", "Kritik Limit", "Operasyonel Alt", "Operasyonel Üst", "İzleme Sıklığı (dk)", "Tanımlı Düzeltici Faaliyet"],
            (8, 16, 30, 7, 16, 12, 12, 12, 60))
    for kkn in plan.values():
        pl.append([kkn.kod, kkn.adim, kkn.parametre, kkn.birim, kkn.limit_metni(), kkn.o_alt and float(kkn.o_alt), kkn.o_ust and float(kkn.o_ust), kkn.siklik, kkn.df])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(plan_yolu: Path, kayit_yolu: Path, cikti: Path, egilim_n: int = 5, tekrar_n: int = 8) -> dict:
    plan = plan_oku(plan_yolu)
    kayit_listesi, hatalar = kayitlari_oku(kayit_yolu)
    s = analiz_et(plan, kayit_listesi, egilim_n, tekrar_n)
    rapor_yaz(cikti, plan, kayit_listesi, s, hatalar)
    return {**s, "plan": plan, "kayitlar": kayit_listesi, "hatalar": hatalar}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="HACCP KKN izleme kayıtlarını kritik limitlerle karşılaştırıp sapmaları ve düzeltici faaliyet ihtiyacını listeler.")
    p.add_argument("--plan", type=Path, default=ORNEK / "kkn_tanimlari.csv", help="HACCP planı: KKN kodu, kritik ve operasyonel limitler, izleme sıklığı")
    p.add_argument("--kayitlar", type=Path, default=ORNEK / "izleme_kayitlari.csv", help="İzleme kayıtları (.xlsx/.csv)")
    p.add_argument("--egilim", type=int, default=5, help="Limite doğru art arda kaç ölçüm eğilim sayılır (varsayılan 5)")
    p.add_argument("--tekrar", type=int, default=8, help="Art arda kaç aynı değer şüpheli sayılır (varsayılan 8)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "haccp_kkn_analizi.xlsx")
    a = p.parse_args(argv)
    for y in (a.plan, a.kayitlar):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        s = calistir(a.plan, a.kayitlar, a.cikti, a.egilim, a.tekrar)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(s['kayitlar'])} kayıt · {len(s['plan'])} KKN · {len(s['sapmalar'])} kritik sapma · {len(s['uyarilar'])} uyarı")
    for x in s["sapmalar"]:
        k = x["kayit"]
        print(f"[X] {x['kkn'].kod} {k.an:%d.%m.%Y %H:%M}: {k.ham} (limit {x['kkn'].limit_metni()})" + ("" if x["df_var"] else " — düzeltici faaliyet kaydı YOK"))
    for u in s["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
