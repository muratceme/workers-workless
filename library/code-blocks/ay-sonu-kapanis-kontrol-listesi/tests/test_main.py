import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def a(sira, n, durum="Başlamadı", bitti=None, onkosul=(), sorumlu="S", kategori="K"):
    return {"sira": str(sira), "adim": f"Adım {sira}", "kategori": kategori, "sorumlu": sorumlu, "n": n, "onkosul": [str(x) for x in onkosul], "durum": durum, "bitti": bitti, "not": ""}


class IsGunuTesti(unittest.TestCase):
    def test_hedef_tarih(self):
        h = lambda n, ay=9, t=(): main.hedef_tarih(2026, ay, n, set(t))   # noqa: E731
        self.assertEqual([h(n) for n in (0, 1, 2, 3, 5, -2, -3)],
                         [date(2026, 9, 30), date(2026, 10, 1), date(2026, 10, 2), date(2026, 10, 5), date(2026, 10, 7), date(2026, 9, 28), date(2026, 9, 25)])
        # Ekim 2026: 31'i Cumartesi → son iş günü 30 Cuma; 29 Ekim tatil
        self.assertEqual((h(0, 10), h(1, 10), h(-1, 10), h(-1, 10, [date(2026, 10, 29)]), h(0, 10, [date(2026, 10, 30)]), h(1, 10, [date(2026, 11, 2)])),
                         (date(2026, 10, 30), date(2026, 11, 2), date(2026, 10, 29), date(2026, 10, 28), date(2026, 10, 29), date(2026, 11, 3)))


class DegerlendirmeTesti(unittest.TestCase):
    def test_izleme_ve_engel(self):
        l = [a(1, 0, "Tamamlandı", date(2026, 9, 30)), a(2, 1, "Tamamlandı", date(2026, 10, 2)), a(3, 1, "Devam ediyor", onkosul=[1, 2]), a(4, 4, onkosul=[3]), a(5, 5, onkosul=[4, 6]),
             a(6, 2, "Uygulanmaz"), a(7, 2, "Tamamlandı", onkosul=[3, 9]), a(8, 3, sorumlu="")]
        r = main.degerlendir(l, 2026, 9, date(2026, 10, 6))
        self.assertEqual([(x["sira"], x["izleme"], x["gecikme"], x["engel"]) for x in r["adimlar"]],
                         [("1", "Tamamlandı", 0, []), ("2", "Gecikmeli tamamlandı", 0, []), ("3", "Gecikti", 5, []), ("4", "Bugün", 0, ["3"]), ("5", "Bekliyor", 0, ["4"]), ("6", "Uygulanmaz", 0, []),
                          ("7", "Tamamlandı", 0, ["3"]), ("8", "Gecikti", 1, [])])
        self.assertEqual((r["kapsam"], r["tamam"], r["oran"], r["geciken"], r["siradaki"]), (7, 3, D("42.86"), 2, ["3", "8"]))
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Geciken adım", "Adım 3"), ("Geciken adım", "Adım 8"), ("Sorumlu yok", "Adım 8"), ("Tamamlanma tarihi yok", "Adım 7"), ("Önkoşul bulunamadı", "Adım 7"), ("Önkoşul tamamlanmadan kapatılmış", "Adım 7")])

    def test_ozet_ve_hata(self):
        l = [a(1, 0, "Tamamlandı", date(2026, 9, 30), kategori="Banka", sorumlu="Ali"), a(2, 1, kategori="Banka", sorumlu="Ali"), a(3, 1, "Tamamlandı", date(2026, 10, 1), kategori="Vergi", sorumlu="Ece"),
             a(4, 9, "Uygulanmaz", kategori="Vergi", sorumlu="Ece"), a(1, 9, kategori="Stok", sorumlu="ali")]
        r = main.degerlendir(l, 2026, 9, date(2026, 10, 6))
        self.assertEqual([(x["ad"], x["toplam"], x["tamam"], x["geciken"], x["oran"]) for x in r["kategoriler"]], [("Banka", 2, 1, 1, D("50.00")), ("Stok", 1, 0, 0, D("0.00")), ("Vergi", 1, 1, 0, D("100.00"))])
        self.assertEqual([(x["ad"], x["toplam"], x["tamam"], x["geciken"]) for x in r["sorumlular"]], [("Ali", 3, 1, 1), ("Ece", 1, 1, 0)])
        self.assertIn("Sıra numarası tekrar ediyor", [u["tur"] for u in r["uyarilar"]])
        self.assertEqual(main.degerlendir([a(1, 0, "Uygulanmaz")], 2026, 9, date(2026, 10, 6))["oran"], None)
        with self.assertRaises(ValueError):
            main.degerlendir([], 2026, 9, date(2026, 10, 6))


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "k.xlsx"
            r = main.calistir(ORNEK / "adimlar.csv", cikti, 2026, 9, date(2026, 10, 6))
            self.assertEqual([(x["sira"], x["izleme"], x["gecikme"]) for x in r["adimlar"] if x["izleme"] not in ("Tamamlandı", "Bekliyor")],
                             [("3", "Gecikmeli tamamlandı", 0), ("5", "Gecikti", 1), ("7", "Gecikti", 4), ("8", "Gecikti", 1), ("9", "Gecikti", 1), ("11", "Bugün", 0), ("12", "Gecikti", 1), ("15", "Uygulanmaz", 0)])
            self.assertEqual((r["kapsam"], r["tamam"], r["oran"], r["geciken"], r["siradaki"], r["adimlar"][12]["engel"], r["adimlar"][13]["engel"]),
                             (14, 6, D("42.86"), 5, ["5", "7", "8", "9", "11", "12"], ["5", "7", "8", "9", "11", "12"], ["13"]))
            self.assertEqual([u["tur"] for u in r["uyarilar"]], ["Geciken adım"] * 5)
            self.assertEqual(load_workbook(cikti).sheetnames, ["Kontrol Listesi", "Kategori Özeti", "Sorumlu Özeti", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, t = Path(tmp) / "a.csv", Path(tmp) / "t.csv"
            y.write_text("No;İş;Grup;Sorumlu;Hedef Gün;Bağımlı;Durum;Tamamlandı\n;Banka;B;Ali;İG+2;;TAMAM;02.10.2026\n;Kur;B;Ali;-1;1;iptal;\n;Stok;S;Ece;x;;;\n;Mizan;R;Ece;3;1, 2;;\n", encoding="utf-8")
            t.write_text("Tarih;Ad\n05.10.2026;Örnek tatil\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["sira"], x["adim"], x["n"], x["onkosul"], x["durum"], x["bitti"]) for x in l], [u["tur"] for u in uy]),
                             ([("1", "Banka", 2, [], "Tamamlandı", date(2026, 10, 2)), ("2", "Kur", -1, ["1"], "Başlamadı", None), ("3", "Mizan", 3, ["1", "2"], "Başlamadı", None)], ["Durum tanınmadı", "Adım okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--adimlar", str(y), "--cikti", c]
            r = main.calistir(y, Path(c), 2026, 9, date(2026, 10, 6), t)
            self.assertEqual([x["hedef"] for x in r["adimlar"]], [date(2026, 10, 2), date(2026, 9, 29), date(2026, 10, 6)])
            self.assertEqual(main.main(arg + ["--donem", "09.2026", "--tarih", "06.10.2026", "--tatiller", str(t)]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual((main.main(arg + ["--donem", "13.2026"]), main.main(arg + ["--tarih", "x"])), (2, 2))
            self.assertEqual((main.main(["--adimlar", str(Path(tmp) / "yok.csv")]), main.main(arg + ["--tatiller", str(Path(tmp) / "yok.csv")])), (1, 1))
            for icerik in ("Adım;Durum\nA;\n", "Adım;Gün\nA;x\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
