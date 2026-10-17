"""
Ay Sonu Kapanış Kontrol Listesi — Gerisi Bende kod bloğu
Muhasebe › Genel Muhasebe Uzmanı

Kapanış adımlarını hedef gün, sorumlu, önkoşul ve durumla izler; gecikenleri ve sıradaki işleri gösterir.
  Hedef tarih        = "Gün" sütunu N ise: N = 0 → dönemin son iş günü;  N > 0 → ondan sonraki N. iş günü;  N < 0 → önceki |N|. iş günü
  İş günü            = hafta içi ve --tatiller dosyasında olmayan gün
  Gecikti            = tamamlanmamış ve hedef tarihi --tarih gününden önce;  gecikme = takvim günü farkı
  Gecikmeli tamam    = tamamlanma tarihi hedef tarihten sonra
  Engelli            = önkoşul adımlarından en az biri tamamlanmamış (Uygulanmaz adım engel sayılmaz)
  Tamamlanma oranı   = tamamlanan ÷ (tüm adımlar − uygulanmaz)
Liste adımları siz verirsiniz; kod kapanışta hangi kayıtların gerektiğine karar vermez ve kayıt atmaz.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
import calendar

SUTUN = {"sira": ("sira", "sira no", "no", "adim no"), "adim": ("adim", "is", "kontrol", "gorev"), "kategori": ("kategori", "alan", "grup"), "sorumlu": ("sorumlu", "sorumlu kisi"),
         "gun": ("gun", "hedef gun", "is gunu", "kapanis gunu"), "onkosul": ("onkosul", "on kosul", "bagimli", "onceki adim"), "durum": ("durum",),
         "bitti": ("tamamlanma tarihi", "tamamlandi", "bitis tarihi"), "not": ("not", "aciklama")}
DURUM = {"tamamlandi": "Tamamlandı", "bitti": "Tamamlandı", "tamam": "Tamamlandı", "devam": "Devam ediyor", "devam ediyor": "Devam ediyor", "basladi": "Devam ediyor", "": "Başlamadı", "baslamadi": "Başlamadı",
         "bekliyor": "Başlamadı", "uygulanmaz": "Uygulanmaz", "yok": "Uygulanmaz", "gecerli degil": "Uygulanmaz"}
KAPALI = ("Tamamlandı", "Uygulanmaz")
GUN = timedelta(days=1)
IZ_RENK = {"Gecikti": KIRMIZI, "Bugün": SARI, "Gecikmeli tamamlandı": MAVI, "Tamamlandı": YESIL}


def is_gunu(g: date, tatiller: set[date]) -> bool:
    return g.weekday() < 5 and g not in tatiller


def hedef_tarih(yil: int, ay: int, n: int, tatiller: set[date]) -> date:
    g = date(yil, ay, calendar.monthrange(yil, ay)[1])
    while not is_gunu(g, tatiller):
        g -= GUN
    adim = GUN if n > 0 else -GUN
    for _ in range(abs(n)):
        g += adim
        while not is_gunu(g, tatiller):
            g += adim
    return g


def tatil_oku(yol: Path) -> set[date]:
    return {t for r in tablo_oku(yol) for t in [tarih(r[0])] if t}


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("adim", "gun")):
        ad = metin(r["adim"])
        if not ad:
            continue
        m = re.search(r"[-−–]?\s*\d+", metin(r["gun"]))
        if not m:
            uy.append(uyari("Yüksek", "Adım okunamadı", f"{ad} (satır {r['_satir']})", f"Gün '{metin(r['gun'])}' okunamadı. 0, 3, -2 veya İG+3 biçiminde yazın. Adım listeye alınmadı."))
            continue
        d = katla(r["durum"])
        if d not in DURUM:
            uy.append(uyari("Orta", "Durum tanınmadı", ad, f"'{metin(r['durum'])}' durumu tanınmadı; adım 'Başlamadı' sayıldı. Tamamlandı, Devam ediyor, Başlamadı veya Uygulanmaz yazın."))
        sira = metin(r["sira"]) or str(len(l) + 1)
        l.append({"sira": sira[:-2] if sira.endswith(".0") else sira, "adim": ad, "kategori": metin(r["kategori"]) or "(kategori yok)", "sorumlu": metin(r["sorumlu"]),
                  "n": int(re.sub(r"\s", "", m.group()).replace("−", "-").replace("–", "-")), "onkosul": [x for x in re.split(r"[,;/\s]+", metin(r["onkosul"])) if x],
                  "durum": DURUM.get(d, "Başlamadı"), "bitti": tarih(r["bitti"]), "not": metin(r["not"])})
    return l, uy


def _ozet(adimlar: list[dict], alan: str) -> list[dict]:
    o = {}
    for a in adimlar:
        if a["durum"] == "Uygulanmaz":
            continue
        x = o.setdefault(katla(a[alan]), {"ad": a[alan] or "(sorumlu yok)", "toplam": 0, "tamam": 0, "geciken": 0})
        x["toplam"] += 1
        x["tamam"] += a["durum"] == "Tamamlandı"
        x["geciken"] += a["izleme"] == "Gecikti"
    return sorted(({**x, "oran": yuvarla(Decimal(100 * x["tamam"]) / x["toplam"])} for x in o.values()), key=lambda x: (-x["geciken"], katla(x["ad"])))


def degerlendir(adimlar: list[dict], yil: int, ay: int, bugun: date, tatiller: set[date] = frozenset()) -> dict:
    if not adimlar:
        raise ValueError("Listede izlenecek kapanış adımı yok.")
    uy, durumlar = [], {}
    for a in adimlar:
        if a["sira"] in durumlar:
            uy.append(uyari("Orta", "Sıra numarası tekrar ediyor", a["sira"], "Aynı sıra numarası birden çok adımda var; önkoşul eşleşmesi son satıra göre yapılır."))
        durumlar[a["sira"]] = a["durum"]
    for a in adimlar:
        a["hedef"] = hedef_tarih(yil, ay, a["n"], tatiller)
        a["gecikme"] = 0
        if a["durum"] == "Uygulanmaz":
            a["izleme"] = "Uygulanmaz"
        elif a["durum"] == "Tamamlandı":
            a["izleme"] = "Gecikmeli tamamlandı" if a["bitti"] and a["bitti"] > a["hedef"] else "Tamamlandı"
            if not a["bitti"]:
                uy.append(uyari("Bilgi", "Tamamlanma tarihi yok", a["adim"], "Adım tamamlandı işaretli ama tarihi yazılmamış; zamanında kapandığı doğrulanamadı."))
        elif a["hedef"] < bugun:
            a["izleme"], a["gecikme"] = "Gecikti", (bugun - a["hedef"]).days
        else:
            a["izleme"] = "Bugün" if a["hedef"] == bugun else "Bekliyor"
        bilinmeyen = [s for s in a["onkosul"] if s not in durumlar or s == a["sira"]]
        if bilinmeyen:
            uy.append(uyari("Orta", "Önkoşul bulunamadı", a["adim"], f"Önkoşul olarak yazılan {', '.join(bilinmeyen)} listede yok veya adımın kendisi; dikkate alınmadı."))
        a["engel"] = [s for s in a["onkosul"] if s not in bilinmeyen and durumlar[s] not in KAPALI]
        if a["engel"] and a["durum"] == "Tamamlandı":
            uy.append(uyari("Orta", "Önkoşul tamamlanmadan kapatılmış", a["adim"], f"Adım tamamlandı işaretli ama önkoşulu olan {', '.join(a['engel'])} numaralı adım(lar) açık. Sırayı ve kaydı kontrol edin."))
        if a["izleme"] == "Gecikti":
            uy.append(uyari("Yüksek", "Geciken adım", a["adim"], f"Hedef {a['hedef']:%d.%m.%Y}, {a['gecikme']} gün gecikmede. Sorumlu: {a['sorumlu'] or 'belirtilmemiş'}."
                            + (f" Önce {', '.join(a['engel'])} numaralı adım(lar) kapanmalı." if a["engel"] else "")))
        if not a["sorumlu"] and a["durum"] != "Uygulanmaz":
            uy.append(uyari("Orta", "Sorumlu yok", a["adim"], "Adıma sorumlu atanmamış."))
    kapsam = [a for a in adimlar if a["durum"] != "Uygulanmaz"]
    tamam = sum(a["durum"] == "Tamamlandı" for a in kapsam)
    return {"adimlar": adimlar, "kapsam": len(kapsam), "tamam": tamam, "oran": yuvarla(Decimal(100 * tamam) / len(kapsam)) if kapsam else None, "geciken": sum(a["izleme"] == "Gecikti" for a in adimlar),
            "siradaki": [a["sira"] for a in adimlar if a["durum"] not in KAPALI and not a["engel"]], "kategoriler": _ozet(adimlar, "kategori"), "sorumlular": _ozet(adimlar, "sorumlu"),
            "donem": f"{ay:02d}.{yil}", "bugun": bugun, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kontrol Listesi"
    _baslik(ws, ["Sıra", "Adım", "Kategori", "Sorumlu", "Hedef Tarih", "Durum", "İzleme", "Gecikme (gün)", "Tamamlanma", "Engelleyen Adımlar", "Not"], (6, 46, 18, 20, 12, 14, 20, 9, 12, 16, 30))
    for a in r["adimlar"]:
        _satir(ws, [a["sira"], a["adim"], a["kategori"], a["sorumlu"], a["hedef"], a["durum"], a["izleme"], a["gecikme"] or None, a["bitti"], ", ".join(a["engel"]), a["not"]], tarih_s=(5, 9),
               dolgu={7: IZ_RENK.get(a["izleme"]), 10: SARI if a["engel"] else None}, sar=(2,))
    ws.append([])
    ws.append([f"Dönem {r['donem']} · durum tarihi {r['bugun']:%d.%m.%Y} · tamamlanan {r['tamam']} / {r['kapsam']}" + (f" (%{r['oran']})" if r["oran"] is not None else "") + f" · geciken {r['geciken']}"])
    for ad, veri, baslik in (("Kategori Özeti", r["kategoriler"], "Kategori"), ("Sorumlu Özeti", r["sorumlular"], "Sorumlu")):
        ws = wb.create_sheet(ad)
        _baslik(ws, [baslik, "Adım", "Tamamlanan", "Geciken", "Tamamlanma %"], (28, 8, 12, 10, 14))
        for x in veri:
            _satir(ws, [x["ad"], x["toplam"], x["tamam"], x["geciken"] or None, x["oran"]], para_s=(5,), dolgu={4: KIRMIZI if x["geciken"] else None})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, yil: int, ay: int, bugun: date, tatil_yol: Path | None = None) -> dict:
    adimlar, uy = oku(yol)
    r = degerlendir(adimlar, yil, ay, bugun, tatil_oku(tatil_yol) if tatil_yol else set())
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Ay sonu kapanış adımlarını hedef gün, sorumlu ve durumla izler.")
    p.add_argument("--adimlar", type=Path, help="Sıra, Adım, Kategori, Sorumlu, Gün, Önkoşul, Durum, Tamamlanma Tarihi, Not")
    p.add_argument("--donem", help="Kapanışı yapılan ay (AA.YYYY); verilmezse geçen ay")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--tatiller", type=Path, help="Resmî tatil günleri (ilk sütun tarih); iş günü hesabında atlanır")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kapanis_kontrol_listesi.xlsx")
    a = p.parse_args(argv)
    ornek = a.adimlar is None
    yol = ORNEK / "adimlar.csv" if ornek else a.adimlar
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 6) if ornek else date.today()
    m = re.fullmatch(r"(\d{1,2})[./-](\d{4})", (a.donem or "").strip())
    if bugun is None or (a.donem and (not m or not 1 <= int(m.group(1)) <= 12)):
        print("[X] --tarih GG.AA.YYYY, --donem AA.YYYY biçiminde olmalı")
        return 2
    onceki = bugun.replace(day=1) - GUN
    ay, yil = (int(m.group(1)), int(m.group(2))) if m else (onceki.month, onceki.year)
    for y in (yol, a.tatiller):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(yol, a.cikti, yil, ay, bugun, a.tatiller)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] Dönem {r['donem']} · {r['tamam']} / {r['kapsam']} adım tamamlandı · {r['geciken']} adım gecikmede · şimdi yapılabilir: {', '.join(r['siradaki']) or '-'}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
