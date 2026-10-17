"""
Talep ve Şikâyet Süre (SLA) Takibi — Gerisi Bende kod bloğu
Müşteri Hizmetleri › Müşteri Temsilcisi

Açık talep ve şikâyetleri hedef cevap ve çözüm sürelerine göre izler; süresi geçen ve yaklaşanları sorumlu bazında listeler.
  Hedef süre      = SLA dosyasındaki kural; arama sırası: konu + öncelik → konu → öncelik → genel satır (konu ve öncelik boş)
                    → komut satırı varsayılanı (--cevap-gun 1, --cozum-gun 5)
  Cevap hedefi    = açılış tarihi + ilk cevap süresi;  Çözüm hedefi = açılış tarihi + çözüm süresi
                    --is-gunu verilirse süre iş günü sayılır (cumartesi ve pazar atlanır; resmi tatiller atlanmaz)
  İlk cevap       = "İlk Cevap Tarihi"; boşsa ve talep kapanmışsa kapanış tarihi
  Açık talep:  Gecikmiş (rapor tarihi > çözüm hedefi) · Bugün son gün (=) · Yaklaşıyor (kalan ≤ --uyari-gun) · Süresi var
  Kapanan talep: Zamanında çözüldü (kapanış ≤ çözüm hedefi) · Geç çözüldü
  Çözüm süresi    = kapanış − açılış (takvim günü);  Zamanında çözüm % = zamanında çözülen ÷ kapanan × 100
Süreler kurumun kendi kuralıdır; kod hiçbir yasal veya sözleşmesel süreyi bilmez, siz verirsiniz. "İptal" durumundaki kayıtlar değerlendirilmez.
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def sayi(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"no": ("talep no", "kayit no", "basvuru no", "sikayet no", "talep numarasi", "no"),
         "acilis": ("acilis tarihi", "talep tarihi", "basvuru tarihi", "kayit tarihi", "tarih"),
         "kanal": ("kanal", "basvuru kanali"), "konu": ("konu", "kategori", "talep turu", "tur"), "oncelik": ("oncelik", "onem"),
         "sorumlu": ("sorumlu", "atanan", "temsilci", "birim"), "cevap": ("ilk cevap tarihi", "ilk cevap", "cevap tarihi", "ilk donus tarihi"),
         "kapanis": ("kapanis tarihi", "cozum tarihi", "kapanis", "sonuclanma tarihi"), "durum": ("durum", "talep durumu")}
SLA = {"konu": SUTUN["konu"], "oncelik": SUTUN["oncelik"], "cevap": ("ilk cevap gun", "ilk cevap suresi", "cevap suresi", "cevap gun", "ilk cevap"),
       "cozum": ("cozum gun", "cozum suresi", "cozum", "kapanis suresi")}
SIRA = ("Gecikmiş", "Bugün son gün", "Yaklaşıyor", "Süresi var", "Geç çözüldü", "Zamanında çözüldü")
ACIK = SIRA[:4]
KAPALI = ("kapandi", "kapali", "kapatildi", "cozuldu", "cozumlendi", "tamamlandi", "sonuclandi")


def gun_ekle(d: date, n: int, is_gunu: bool) -> date:
    if not is_gunu:
        return d + timedelta(days=n)
    while n > 0:
        d += timedelta(days=1)
        n -= d.weekday() < 5
    return d


def oku(yol: Path, sla_yol: Path | None = None) -> tuple[list[dict], dict, list[dict]]:
    uy, l, sla, gorulen = [], [], {}, set()
    for r in kayitlar(yol, SUTUN, ("no", "acilis")):
        no = metin(r["no"])
        if not no:
            continue
        kim = f"{no} (satır {r['_satir']})"
        ac = tarih(r["acilis"])
        if ac is None:
            uy.append(uyari("Yüksek", "Açılış tarihi okunamadı", kim, f"'{metin(r['acilis'])}'. Talep değerlendirilmedi."))
            continue
        if katla(no) in gorulen:
            uy.append(uyari("Orta", "Mükerrer talep no", kim, "Aynı talep numarası birden fazla satırda; hepsi ayrı değerlendirildi. Kayıt iki kez aktarılmış olabilir."))
        gorulen.add(katla(no))
        x = {"no": no, "acilis": ac, "durum": metin(r["durum"])}
        for a in ("kanal", "konu", "oncelik", "sorumlu"):
            x[a] = metin(r[a]) or "Belirtilmemiş"
        for a, ad in (("cevap", "İlk cevap tarihi"), ("kapanis", "Kapanış tarihi")):
            x[a] = tarih(r[a])
            if x[a] is None and metin(r[a]):
                uy.append(uyari("Yüksek", f"{ad} okunamadı", kim, f"'{metin(r[a])}'. Boş sayıldı."))
            elif x[a] and x[a] < ac:
                uy.append(uyari("Yüksek", f"{ad} açılıştan önce", kim, f"{ad} {x[a]:%d.%m.%Y}, açılış {ac:%d.%m.%Y}. Tarih yok sayıldı; kaydı düzeltin."))
                x[a] = None
        if x["cevap"] and x["kapanis"] and x["cevap"] > x["kapanis"]:
            uy.append(uyari("Orta", "İlk cevap kapanıştan sonra", kim, f"İlk cevap {x['cevap']:%d.%m.%Y}, kapanış {x['kapanis']:%d.%m.%Y}. Tarihleri kontrol edin."))
        if not x["kapanis"] and any(k in katla(x["durum"]) for k in KAPALI):
            uy.append(uyari("Orta", "Kapalı görünüyor, kapanış tarihi yok", kim, f"Durum '{x['durum']}' ama kapanış tarihi boş; talep açık sayıldı. Kapanış tarihini girin."))
        l.append(x)
    if sla_yol:
        for r in kayitlar(sla_yol, SLA, ("cozum",)):
            k, o = katla(r["konu"]), katla(r["oncelik"])
            ad = " / ".join(v for v in (metin(r["konu"]), metin(r["oncelik"])) if v) or "Genel kural"
            kural, hatali = {"ad": ad}, False
            for a in ("cevap", "cozum"):
                v = sayi(r[a])
                if v is None and not metin(r[a]):
                    kural[a] = None
                elif v is None or v < 0 or v != int(v):
                    hatali = True
                else:
                    kural[a] = int(v)
            if hatali:
                uy.append(uyari("Yüksek", "SLA süresi okunamadı", ad, f"{sla_yol.name} satır {r['_satir']}: süreler sıfır veya pozitif tam gün olmalı. Kural kullanılmadı."))
            elif (k, o) in sla:
                uy.append(uyari("Orta", "Mükerrer SLA kuralı", ad, f"{sla_yol.name} satır {r['_satir']}: aynı konu/öncelik için ikinci kural; ilki kullanıldı."))
            else:
                sla[(k, o)] = kural
    return l, sla, uy


def hedef(t: dict, alan: str, sla: dict, varsayilan: int) -> tuple[int, str]:
    k, o = katla(t["konu"]), katla(t["oncelik"])
    for anahtar in ((k, o), (k, ""), ("", o), ("", "")):
        kural = sla.get(anahtar)
        if kural and kural.get(alan) is not None:
            return kural[alan], kural["ad"]
    return varsayilan, "Varsayılan"


def degerlendir(talepler: list[dict], rapor: date, sla: dict | None = None, cevap_gun: int = 1, cozum_gun: int = 5, uyari_gun: int = 1, is_gunu: bool = False) -> dict:
    sla = sla or {}
    iptal = [t for t in talepler if "iptal" in katla(t["durum"])]
    l = [t for t in talepler if "iptal" not in katla(t["durum"])]
    if not l:
        raise ValueError("Değerlendirilecek talep yok.")
    varsayilan = set()
    for t in l:
        cg, _ = hedef(t, "cevap", sla, cevap_gun)
        zg, t["kural"] = hedef(t, "cozum", sla, cozum_gun)
        if t["kural"] == "Varsayılan":
            varsayilan.add(f"{t['konu']} / {t['oncelik']}")
        t["hedef_cevap"], t["hedef_cozum"] = gun_ekle(t["acilis"], cg, is_gunu), gun_ekle(t["acilis"], zg, is_gunu)
        ilk = t["cevap"] or t["kapanis"]
        if ilk:
            t["cevap_sonuc"] = "Zamanında" if ilk <= t["hedef_cevap"] else "Geç"
        else:
            t["cevap_sonuc"] = "Gecikmiş" if rapor > t["hedef_cevap"] else "Bekliyor"
        t["gecikme"] = t["kalan"] = t["sure"] = t["yas"] = None
        if t["kapanis"]:
            t["sonuc"] = "Zamanında çözüldü" if t["kapanis"] <= t["hedef_cozum"] else "Geç çözüldü"
            t["gecikme"] = (t["kapanis"] - t["hedef_cozum"]).days if t["kapanis"] > t["hedef_cozum"] else None
            t["sure"] = (t["kapanis"] - t["acilis"]).days
        else:
            t["kalan"], t["yas"] = (t["hedef_cozum"] - rapor).days, (rapor - t["acilis"]).days
            t["sonuc"] = ("Gecikmiş" if t["kalan"] < 0 else "Bugün son gün" if t["kalan"] == 0 else "Yaklaşıyor" if t["kalan"] <= uyari_gun else "Süresi var")
            t["gecikme"] = -t["kalan"] if t["kalan"] < 0 else None

    def ozet(alan: str) -> list[dict]:
        g = {}
        for t in l:
            x = g.setdefault(katla(t[alan]), {alan: t[alan], "toplam": 0, "sureler": [], "gunler": [], "cevap_gec": 0} | {s: 0 for s in SIRA})
            x["toplam"] += 1
            x[t["sonuc"]] += 1
            x["cevap_gec"] += t["cevap_sonuc"] in ("Geç", "Gecikmiş")
            if t["sure"] is not None:
                x["sureler"].append(t["sure"])
            if t["sonuc"] == "Gecikmiş":
                x["gunler"].append(t["gecikme"])
        for x in g.values():
            kapanan = len(x["sureler"])
            x["kapanan"] = kapanan
            x["zamaninda_oran"] = yuvarla(Decimal(x["Zamanında çözüldü"]) / kapanan * 100) if kapanan else None
            x["ortalama_sure"] = yuvarla(Decimal(sum(x["sureler"])) / kapanan) if kapanan else None
            x["en_uzun"] = max(x["gunler"], default=None)
        return sorted(g.values(), key=lambda x: (-x["Gecikmiş"], -(x["en_uzun"] or 0), -x["Bugün son gün"], x[alan]))

    sureler = [t["sure"] for t in l if t["sure"] is not None]
    kapanan = len(sureler)
    return {"talepler": sorted(l, key=lambda t: (SIRA.index(t["sonuc"]), -(t["gecikme"] or 0), t["hedef_cozum"], t["no"])),
            "sorumlu": ozet("sorumlu"), "konu": ozet("konu"), "rapor": rapor, "iptal": len(iptal),
            "say": {s: sum(t["sonuc"] == s for t in l) for s in SIRA}, "cevap_gecikmis": sum(t["cevap_sonuc"] == "Gecikmiş" for t in l),
            "zamaninda_oran": yuvarla(Decimal(sum(t["sonuc"] == "Zamanında çözüldü" for t in l)) / kapanan * 100) if kapanan else None,
            "ortalama_sure": yuvarla(Decimal(sum(sureler)) / kapanan) if kapanan else None,
            "ileri_tarihli": [t["no"] for t in l if t["acilis"] > rapor],
            "ayar": (cevap_gun, cozum_gun, uyari_gun, is_gunu), "varsayilan": sorted(varsayilan) if sla else []}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict], sla_var: bool) -> None:
    wb = Workbook()
    renk = {"Gecikmiş": KIRMIZI, "Bugün son gün": SARI, "Yaklaşıyor": SARI, "Geç çözüldü": MAVI, "Zamanında çözüldü": YESIL}
    crenk = {"Gecikmiş": KIRMIZI, "Geç": MAVI}
    cg, zg, ug, ig = r["ayar"]
    ws = wb.active
    ws.title = "Açık Talepler"
    _baslik(ws, ["Durum", "Talep No", "Konu", "Öncelik", "Sorumlu", "Kanal", "Açılış", "Yaş (gün)", "İlk Cevap", "Cevap Hedefi", "İlk Cevap Durumu",
                 "Çözüm Hedefi", "Kalan Gün", "Gecikme (gün)", "Uygulanan Kural", "Talep Durumu", "Not (siz doldurun)"],
            (14, 14, 16, 10, 16, 12, 12, 8, 12, 12, 12, 12, 8, 9, 18, 14, 30))
    for t in r["talepler"]:
        if t["sonuc"] in ACIK:
            _satir(ws, [t["sonuc"], t["no"], t["konu"], t["oncelik"], t["sorumlu"], t["kanal"], t["acilis"], t["yas"], t["cevap"], t["hedef_cevap"], t["cevap_sonuc"],
                        t["hedef_cozum"], t["kalan"], t["gecikme"], t["kural"], t["durum"], ""], tarih_s=(7, 9, 10, 12), dolgu={1: renk.get(t["sonuc"]), 11: crenk.get(t["cevap_sonuc"])})
    ws.append([])
    ws.append([f"Rapor tarihi {r['rapor']:%d.%m.%Y}. Süreler {'SLA dosyasından; kuralı olmayanlarda ' if sla_var else ''}varsayılan ilk cevap {cg}, çözüm {zg} gün "
               f"({'iş günü' if ig else 'takvim günü'}). 'Yaklaşıyor' = kalan ≤ {ug} gün. Yaş ve gecikme takvim günüdür. İptal edilen {r['iptal']} kayıt değerlendirilmedi."])
    for ad, alan, bas in (("Sorumlu Özeti", "sorumlu", "Sorumlu"), ("Konu Özeti", "konu", "Konu")):
        ws = wb.create_sheet(ad)
        _baslik(ws, [bas, "Talep", *ACIK, "Kapanan", "Zamanında Çözüm %", "Ort. Çözüm Süresi (gün)", "İlk Cevabı Geç", "En Uzun Gecikme (gün)"], (22, 8, 10, 10, 11, 10, 9, 12, 12, 10, 12))
        for x in r[alan]:
            _satir(ws, [x[alan], x["toplam"], *[x[s] or None for s in ACIK], x["kapanan"] or None, x["zamaninda_oran"], x["ortalama_sure"], x["cevap_gec"] or None, x["en_uzun"]],
                   para_s=(8, 9), dolgu={3: KIRMIZI if x["Gecikmiş"] else None})
    ws = wb.create_sheet("Kapananlar")
    _baslik(ws, ["Sonuç", "Talep No", "Konu", "Öncelik", "Sorumlu", "Açılış", "İlk Cevap", "Kapanış", "Çözüm Hedefi", "Çözüm Süresi (gün)", "Gecikme (gün)", "İlk Cevap Durumu", "Uygulanan Kural"],
            (18, 14, 16, 10, 16, 12, 12, 12, 12, 10, 9, 12, 18))
    for t in r["talepler"]:
        if t["sonuc"] not in ACIK:
            _satir(ws, [t["sonuc"], t["no"], t["konu"], t["oncelik"], t["sorumlu"], t["acilis"], t["cevap"], t["kapanis"], t["hedef_cozum"], t["sure"], t["gecikme"], t["cevap_sonuc"], t["kural"]],
                   tarih_s=(6, 7, 8, 9), dolgu={1: renk.get(t["sonuc"]), 12: crenk.get(t["cevap_sonuc"])})
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, rapor: date, sla_yol: Path | None = None, cevap_gun: int = 1, cozum_gun: int = 5, uyari_gun: int = 1, is_gunu: bool = False) -> dict:
    talepler, sla, uy = oku(yol, sla_yol)
    r = degerlendir(talepler, rapor, sla, cevap_gun, cozum_gun, uyari_gun, is_gunu)
    for k in r["varsayilan"]:
        uy.append(uyari("Orta", "SLA kuralı tanımlı değil", k, f"SLA dosyasında bu konu/öncelik için kural yok; varsayılan ilk cevap {cevap_gun}, çözüm {cozum_gun} gün kullanıldı."))
    for no in r["ileri_tarihli"]:
        uy.append(uyari("Orta", "Açılış rapor tarihinden sonra", no, f"Açılış tarihi rapor tarihinden ({rapor:%d.%m.%Y}) sonra. Rapor tarihini veya kaydı kontrol edin."))
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy, bool(sla_yol))
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Açık talep ve şikâyetleri hedef cevap ve çözüm sürelerine göre izler.")
    p.add_argument("--talepler", type=Path, default=ORNEK / "talepler.csv",
                   help="Talep No, Açılış Tarihi (zorunlu); Kanal, Konu, Öncelik, Sorumlu, İlk Cevap Tarihi, Kapanış Tarihi, Durum")
    p.add_argument("--sla", type=Path, help="İsteğe bağlı SLA kuralları: Konu, Öncelik, İlk Cevap (gün), Çözüm (gün)")
    p.add_argument("--cevap-gun", default="1", help="Kuralı olmayan talepler için ilk cevap süresi (varsayılan 1 gün)")
    p.add_argument("--cozum-gun", default="5", help="Kuralı olmayan talepler için çözüm süresi (varsayılan 5 gün)")
    p.add_argument("--uyari-gun", default="1", help="Çözüm hedefine bu kadar gün veya daha az kalan açık talepler 'Yaklaşıyor' olur (varsayılan 1)")
    p.add_argument("--is-gunu", action="store_true", help="Süreleri iş günü say (cumartesi ve pazar atlanır; resmi tatiller atlanmaz)")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 15.10.2026)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "talep_sla_takibi.xlsx")
    a = p.parse_args(argv)
    ornek = a.talepler == ORNEK / "talepler.csv"
    sla_yol = a.sla or (ORNEK / "sla_kurallari.csv" if ornek else None)
    rapor = tarih(a.tarih) if a.tarih else (date(2026, 10, 15) if ornek else date.today())
    gunler = [sayi(v) for v in (a.cevap_gun, a.cozum_gun, a.uyari_gun)]
    if rapor is None or any(g is None or g < 0 or g != int(g) for g in gunler):
        print("[X] --tarih GG.AA.YYYY biçiminde; --cevap-gun, --cozum-gun ve --uyari-gun sıfır veya pozitif tam sayı olmalı")
        return 2
    for y in (a.talepler, sla_yol):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.talepler, a.cikti, rapor, sla_yol, *(int(g) for g in gunler), a.is_gunu)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    s = r["say"]
    print(f"[OK] {len(r['talepler'])} talep · açık: {s['Gecikmiş']} gecikmiş, {s['Bugün son gün']} bugün son gün, {s['Yaklaşıyor']} yaklaşıyor, {s['Süresi var']} süresi var · "
          f"kapanan: {s['Zamanında çözüldü']} zamanında, {s['Geç çözüldü']} geç · ilk cevabı gecikmiş açık talep: {r['cevap_gecikmis']}")
    for t in r["talepler"]:
        if t["sonuc"] == "Gecikmiş":
            print(f"[!] {t['no']} · {t['konu']} · {t['sorumlu']}: çözüm hedefi {t['hedef_cozum']:%d.%m.%Y}, {t['gecikme']} gün gecikmiş")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
