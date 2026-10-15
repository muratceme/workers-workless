import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
RAPOR = date(2026, 10, 15)          # perşembe


def tl(no="T1", acilis=-3, cevap=None, kapanis=None, konu="K", oncelik="Normal", sorumlu="S", durum=""):
    t = lambda g: None if g is None else RAPOR + timedelta(days=g)   # noqa: E731
    return {"no": no, "acilis": t(acilis), "cevap": t(cevap), "kapanis": t(kapanis), "konu": konu, "oncelik": oncelik, "sorumlu": sorumlu, "kanal": "E-posta", "durum": durum}


def tek(**kw):
    ayar = {k: kw.pop(k) for k in ("sla", "cevap_gun", "cozum_gun", "uyari_gun", "is_gunu") if k in kw}
    return main.degerlendir([tl(**kw)], RAPOR, **ayar)["talepler"][0]


class DurumTesti(unittest.TestCase):
    def test_acik_talep_sinirlari(self):
        # varsayılan çözüm 5 gün, uyarı 1 gün
        self.assertEqual([(x["sonuc"], x["kalan"], x["gecikme"], x["yas"]) for x in (tek(acilis=g) for g in (-7, -6, -5, -4, -3))],
                         [("Gecikmiş", -2, 2, 7), ("Gecikmiş", -1, 1, 6), ("Bugün son gün", 0, None, 5), ("Yaklaşıyor", 1, None, 4), ("Süresi var", 2, None, 3)])
        self.assertEqual(tek(acilis=-3, uyari_gun=2)["sonuc"], "Yaklaşıyor")
        self.assertEqual(tek(acilis=-4, uyari_gun=0)["sonuc"], "Süresi var")
        self.assertEqual(tek(acilis=0, cozum_gun=0)["sonuc"], "Bugün son gün")

    def test_ilk_cevap(self):
        self.assertEqual(tek(acilis=-1)["cevap_sonuc"], "Bekliyor")                     # hedef bugün
        self.assertEqual(tek(acilis=-2)["cevap_sonuc"], "Gecikmiş")
        self.assertEqual(tek(acilis=-3, cevap=-2)["cevap_sonuc"], "Zamanında")
        self.assertEqual(tek(acilis=-3, cevap=-1)["cevap_sonuc"], "Geç")
        self.assertEqual(tek(acilis=-5, kapanis=-4)["cevap_sonuc"], "Zamanında")          # cevap yoksa kapanış ilk cevap sayılır
        self.assertEqual(tek(acilis=-5, kapanis=-3)["cevap_sonuc"], "Geç")

    def test_kapanan_talep(self):
        x = tek(acilis=-10, kapanis=-5)
        self.assertEqual((x["sonuc"], x["sure"], x["gecikme"], x["kalan"]), ("Zamanında çözüldü", 5, None, None))
        x = tek(acilis=-10, kapanis=-2)
        self.assertEqual((x["sonuc"], x["sure"], x["gecikme"]), ("Geç çözüldü", 8, 3))

    def test_sla_kural_sirasi(self):
        sla = {("", ""): {"ad": "Genel kural", "cevap": 2, "cozum": 9}, ("", "yuksek"): {"ad": "Yüksek", "cevap": 0, "cozum": 2},
               ("fatura", ""): {"ad": "Fatura", "cevap": None, "cozum": 3}, ("fatura", "yuksek"): {"ad": "Fatura / Yüksek", "cevap": 0, "cozum": 1}}
        kural = lambda konu, onc: (main.hedef({"konu": konu, "oncelik": onc}, "cozum", sla, 5), main.hedef({"konu": konu, "oncelik": onc}, "cevap", sla, 1))  # noqa: E731
        self.assertEqual(kural("Fatura", "Yüksek"), ((1, "Fatura / Yüksek"), (0, "Fatura / Yüksek")))
        self.assertEqual(kural("FATURA", "Normal"), ((3, "Fatura"), (2, "Genel kural")))      # konu kuralında cevap boş: sıradaki kurala geçer
        self.assertEqual(kural("Ürün", "yüksek"), ((2, "Yüksek"), (0, "Yüksek")))
        self.assertEqual(kural("Ürün", "Düşük"), ((9, "Genel kural"), (2, "Genel kural")))
        self.assertEqual(main.hedef({"konu": "Ürün", "oncelik": "Düşük"}, "cozum", {}, 5), (5, "Varsayılan"))

    def test_is_gunu(self):
        cuma = date(2026, 10, 9)
        self.assertEqual([main.gun_ekle(cuma, n, True) for n in (0, 1, 5)], [cuma, date(2026, 10, 12), date(2026, 10, 16)])
        self.assertEqual(main.gun_ekle(date(2026, 10, 10), 1, True), date(2026, 10, 12))      # cumartesi + 1 iş günü = pazartesi
        # 08.10 perşembe + 5 iş günü = 15.10; takvim günüyle 13.10
        self.assertEqual((tek(acilis=-7)["sonuc"], tek(acilis=-7, is_gunu=True)["sonuc"]), ("Gecikmiş", "Bugün son gün"))

    def test_ozet_siralama_ve_iptal(self):
        l = [tl("A", -7, sorumlu="Ali"), tl("B", -9, sorumlu="Can"), tl("C", -5, sorumlu="Ali"), tl("D", -10, kapanis=-6, sorumlu="Can"),
             tl("E", -10, cevap=-7, kapanis=-2, sorumlu="Ali"), tl("F", -9, durum="İptal")]
        r = main.degerlendir(l, RAPOR)
        self.assertEqual([t["no"] for t in r["talepler"]], ["B", "A", "C", "E", "D"])
        self.assertEqual((r["say"]["Gecikmiş"], r["say"]["Bugün son gün"], r["iptal"], r["zamaninda_oran"], r["ortalama_sure"], r["cevap_gecikmis"]),
                         (2, 1, 1, D("50.00"), D("6.00"), 3))
        self.assertEqual([(x["sorumlu"], x["toplam"], x["Gecikmiş"], x["kapanan"], x["zamaninda_oran"], x["ortalama_sure"], x["cevap_gec"], x["en_uzun"]) for x in r["sorumlu"]],
                         [("Can", 2, 1, 1, D("100.00"), D("4.00"), 2, 4), ("Ali", 3, 1, 1, D("0.00"), D("8.00"), 3, 2)])
        with self.assertRaises(ValueError):
            main.degerlendir([tl(durum="iptal edildi")], RAPOR)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "s.xlsx"
            r = main.calistir(ORNEK / "talepler.csv", cikti, RAPOR, ORNEK / "sla_kurallari.csv")
            self.assertEqual([(t["no"], t["sonuc"], t["gecikme"], t["kural"]) for t in r["talepler"]],
                             [("T-209", "Gecikmiş", 10, "Fatura"), ("T-201", "Gecikmiş", 2, "Normal"), ("T-204", "Bugün son gün", None, "Fatura"),
                              ("T-211", "Bugün son gün", None, "Normal"), ("T-202", "Yaklaşıyor", None, "Yüksek"), ("T-203", "Yaklaşıyor", None, "Fatura"),
                              ("T-213", "Süresi var", None, "Varsayılan"), ("T-205", "Süresi var", None, "Düşük"), ("T-207", "Geç çözüldü", 3, "Normal"),
                              ("T-206", "Zamanında çözüldü", None, "İade / Yüksek"), ("T-208", "Zamanında çözüldü", None, "Normal"), ("T-212", "Zamanında çözüldü", None, "Normal")])
            self.assertEqual({x["sorumlu"]: (x["toplam"], x["Gecikmiş"], x["en_uzun"], x["kapanan"], x["zamaninda_oran"], x["ortalama_sure"], x["cevap_gec"]) for x in r["sorumlu"]},
                             {"Zeynep": (4, 1, 10, 1, D("100.00"), D("1.00"), 1), "Ayşe": (4, 1, 2, 1, D("0.00"), D("8.00"), 1),
                              "Mehmet": (4, 0, None, 2, D("100.00"), D("3.00"), 1)})
            self.assertEqual([x["sorumlu"] for x in r["sorumlu"]], ["Zeynep", "Ayşe", "Mehmet"])
            self.assertEqual((r["iptal"], r["zamaninda_oran"], r["ortalama_sure"], r["cevap_gecikmis"]), (1, D("75.00"), D("3.75"), 2))
            self.assertEqual([(u["tur"], u["kim"]) for u in r["uyarilar"]], [("Kapalı görünüyor, kapanış tarihi yok", "T-211 (satır 12)"), ("SLA kuralı tanımlı değil", "Ürün / Acil")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Açık Talepler", "Sorumlu Özeti", "Konu Özeti", "Kapananlar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, s = Path(tmp) / "t.csv", Path(tmp) / "s.csv"
            y.write_text("Kayıt No;Başvuru Tarihi;Kategori;İlk Cevap;Kapanış Tarihi;Durum\n"
                         "A;10.10.2026;X;yarın;;\nA;10.10.2026;X;09.10.2026;;\nB;dün;X;;;\nC;2026-10-11 14:30:00;Y;14.10.2026;12.10.2026;Kapandı\nD;12.10.2026;Y;;;\n",
                         encoding="utf-8")
            s.write_text("Konu;Öncelik;İlk Cevap (gün);Çözüm (gün)\nX;;1;3\nX;;2;4\nY;;hızlı;2\n", encoding="utf-8")
            l, sla, uy = main.oku(y, s)
            self.assertEqual([(t["no"], t["acilis"], t["cevap"], t["kapanis"], t["sorumlu"]) for t in l],
                             [("A", date(2026, 10, 10), None, None, "Belirtilmemiş"), ("A", date(2026, 10, 10), None, None, "Belirtilmemiş"),
                              ("C", date(2026, 10, 11), date(2026, 10, 14), date(2026, 10, 12), "Belirtilmemiş"), ("D", date(2026, 10, 12), None, None, "Belirtilmemiş")])
            self.assertEqual(sla, {("x", ""): {"ad": "X", "cevap": 1, "cozum": 3}})
            self.assertEqual(sorted(u["tur"] for u in uy), ["Açılış tarihi okunamadı", "Mükerrer SLA kuralı", "Mükerrer talep no", "SLA süresi okunamadı",
                                                            "İlk cevap kapanıştan sonra", "İlk cevap tarihi açılıştan önce", "İlk cevap tarihi okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            r = main.calistir(y, Path(c), RAPOR, s)
            self.assertEqual([u["kim"] for u in r["uyarilar"] if u["tur"] == "SLA kuralı tanımlı değil"], ["Y / Belirtilmemiş"])
            self.assertEqual([u["kim"] for u in main.calistir(y, Path(c), date(2026, 10, 11), s)["uyarilar"] if u["tur"] == "Açılış rapor tarihinden sonra"], ["D"])
            arg = ["--talepler", str(y), "--cikti", c, "--tarih", "15.10.2026"]
            self.assertEqual(main.main(arg + ["--sla", str(s), "--cozum-gun", "3", "--uyari-gun", "2", "--is-gunu"]), 0)
            self.assertEqual(main.main(arg + ["--cevap-gun", "1,5"]), 2)
            self.assertEqual(main.main(arg + ["--uyari-gun", "-1"]), 2)
            self.assertEqual(main.main(arg[:-1] + ["yarın"]), 2)
            self.assertEqual(main.main(["--talepler", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Talep No;Açılış Tarihi\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
