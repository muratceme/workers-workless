"""
Reklam Kampanyası Performans Raporu — Workers / Workless kod bloğu
Pazarlama › Dijital Pazarlama Uzmanı

Reklam platformu dışa aktarımlarını (Google Ads, Meta Ads, TikTok Ads, Microsoft Ads … — Türkçe veya İngilizce arayüz)
tek tabloda birleştirir ve kampanya bazında raporlar:
  - Harcama, gösterim, tıklama, dönüşüm, dönüşüm değeri; CTR, TBM (CPC), BGBM (CPM), dönüşüm oranı, dönüşüm başı
    maliyet (CPA) ve ROAS.
  - Platform, kampanya ve hafta kırılımları; hedef ROAS / hedef CPA'ya göre işaretler: harcayıp dönüşüm getirmeyen,
    hedefin altında kalan ve ölçeklenebilecek kampanyalar.
  - Bütçe verilirse dönem bütçesine göre harcama temposu (pacing).
Sütunlar başlık adlarından otomatik tanınır; tanınmayanlar için --esleme ile elle eşleştirme yapılabilir.
İnternete bağlanmaz.

Kullanım:
    python main.py                                              # örnek dışa aktarımlarla dener
    python main.py --dosya google_ads.csv --dosya meta_ads.csv --hedef-roas 4 --hedef-cpa 150
    python main.py --dosya tiktok.xlsx --platform TikTok --butce butceler.xlsx
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import OrderedDict, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent

# Alan → başlıkta aranacak ifadeler (katlanmış, ilk eşleşen kazanır; daha özel ifadeler önce)
ALANLAR = OrderedDict([
    ("kampanya", ["kampanya adi", "kampanya", "campaign name", "campaign"]),
    ("tarih", ["gun", "tarih", "day", "date", "reporting starts", "raporlama baslangici"]),
    ("gosterim", ["gosterimler", "gosterim", "impressions", "impr"]),
    ("tiklama", ["baglanti tiklamalari", "link clicks", "tiklamalar", "tiklama", "clicks all", "clicks"]),
    ("harcama", ["harcanan tutar", "amount spent", "maliyet", "harcama", "cost", "spend"]),
    ("donusum_degeri", ["satin alma donusum degeri", "purchases conversion value", "donusum degeri", "conv value",
                        "conversion value", "total conversion value", "purchase value"]),
    ("donusum", ["satin almalar", "purchases", "donusumler", "donusum", "conversions", "results", "sonuclar"]),
])


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def katla(s) -> str:
    s = kucuk(s)
    for a, b in zip("ıöüşğç", "iousgc"):
        s = s.replace(a, b)
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def sayi(x) -> float:
    if x in (None, "", "--", "-"):
        return 0.0
    if isinstance(x, (int, float)):
        return float(x)
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        # 1,234 (binlik) ile 12,5 (ondalık) ayrımı: virgülden sonra tam 3 hane ve başka ayırıcı yoksa binlik
        s = s.replace(",", "") if re.fullmatch(r"-?\d{1,3}(,\d{3})+", s) else s.replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")                                  # 2.500 / 1.234.567 → Türkçe binlik
    try:
        return float(s)
    except ValueError:
        return 0.0


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for b in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(s[:10], b).date()
        except ValueError:
            continue
    return None


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        ham = yol.read_bytes()
        metin = None
        for kod in ("utf-8-sig", "utf-16", "cp1254"):
            try:
                metin = ham.decode(kod)
                if kod == "utf-16" and "\x00" in metin:
                    continue
                break
            except UnicodeDecodeError:
                continue
        satir_l = metin.splitlines()
        ilk = "\n".join(satir_l[:10])
        satirlar = list(csv.reader(satir_l, delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def platform_tahmini(yol: Path, baslik: list[str]) -> str:
    ad = katla(yol.stem)
    b = " ".join(katla(x) for x in baslik)
    if "meta" in ad or "facebook" in ad or "amount spent" in b or "harcanan tutar" in b:
        return "Meta"
    if "tiktok" in ad:
        return "TikTok"
    if "bing" in ad or "microsoft" in ad:
        return "Microsoft"
    if "google" in ad or "impr" in b or "conv value" in b:
        return "Google"
    return yol.stem


def dosya_oku(yol: Path, platform: str | None, esleme: dict[str, str]) -> tuple[list[dict], dict, str]:
    s = tablo_oku(yol)
    # Üstte rapor adı/tarih aralığı satırları olabilir: kampanya sütunu içeren ilk satır başlıktır
    bi = next((i for i, r in enumerate(s[:10]) if sum(1 for c in r if c not in (None, "")) >= 3
               and any(katla(c).startswith(("kampanya", "campaign")) for c in r)), 0)
    b = [katla(x) for x in s[bi]]
    i: dict[str, int | None] = {}
    for alan, ifadeler in ALANLAR.items():
        if alan in esleme:
            hedef = katla(esleme[alan])
            i[alan] = b.index(hedef) if hedef in b else None
            continue
        i[alan] = None
        for ifade in ifadeler:
            aday = [j for j, x in enumerate(b) if (x == ifade or x.startswith(ifade + " ")) and j not in i.values()]
            if aday:
                i[alan] = aday[0]
                break
    if i["kampanya"] is None or i["harcama"] is None:
        raise SystemExit(f"{yol.name}: kampanya ve harcama sütunları tanınamadı. Başlıklar: {s[bi]} — --esleme harcama='…' ile belirtin.")
    para = re.search(r"\(([A-Z]{3})\)", " ".join(str(x) for x in s[bi]))
    plat = platform or platform_tahmini(yol, s[bi])
    satirlar = []
    for r in s[bi + 1:]:
        al = lambda k: r[i[k]] if i[k] is not None and i[k] < len(r) else None  # noqa: E731
        kamp = str(al("kampanya") or "").strip()
        if not kamp or katla(kamp).startswith(("toplam", "total")):
            continue
        satirlar.append({"platform": plat, "kampanya": kamp, "tarih": tarih(al("tarih")),
                         **{k: sayi(al(k)) for k in ("gosterim", "tiklama", "harcama", "donusum", "donusum_degeri")}})
    bulunan = {k: (s[bi][v] if v is not None else None) for k, v in i.items()}
    return satirlar, bulunan, para.group(1) if para else ""


# ----------------------------------------------------------------------------
# Hesap
# ----------------------------------------------------------------------------

def olcutler(x: dict) -> dict:
    g, t, h, d, dd = x["gosterim"], x["tiklama"], x["harcama"], x["donusum"], x["donusum_degeri"]
    return {**x, "ctr": t / g if g else None, "cpc": h / t if t else None, "cpm": h / g * 1000 if g else None,
            "cvr": d / t if t else None, "cpa": h / d if d else None, "roas": dd / h if h else None}


def topla(satirlar, anahtar_fn) -> "OrderedDict":
    gr = defaultdict(lambda: {"gosterim": 0.0, "tiklama": 0.0, "harcama": 0.0, "donusum": 0.0, "donusum_degeri": 0.0})
    for s in satirlar:
        x = gr[anahtar_fn(s)]
        for k in ("gosterim", "tiklama", "harcama", "donusum", "donusum_degeri"):
            x[k] += s[k]
    return OrderedDict((k, olcutler(v)) for k, v in sorted(gr.items(), key=lambda i: -i[1]["harcama"]))


def isaretle(v: dict, hedef_roas: float | None, hedef_cpa: float | None, toplam_harcama: float) -> str:
    notlar = []
    pay = v["harcama"] / toplam_harcama if toplam_harcama else 0
    if v["harcama"] > 0 and v["donusum"] == 0 and pay >= 0.02:
        notlar.append("harcıyor ama dönüşüm yok")
    if hedef_roas and v["roas"] is not None and v["donusum_degeri"]:
        if v["roas"] < hedef_roas * 0.7:
            notlar.append(f"ROAS hedefin çok altında ({v['roas']:.2f} / {hedef_roas:g})".replace(".", ","))
        elif v["roas"] >= hedef_roas * 1.3 and pay >= 0.05:
            notlar.append("hedefin belirgin üstünde — bütçe artırılabilir")
    if hedef_cpa and v["cpa"] is not None and v["cpa"] > hedef_cpa * 1.3:
        notlar.append(f"CPA hedefin üstünde ({v['cpa']:,.0f} / {hedef_cpa:,.0f})".replace(",", "."))
    if v["gosterim"] >= 1000 and v["ctr"] is not None and v["ctr"] < 0.005:
        notlar.append("tıklama oranı çok düşük (< %0,5)")
    return "; ".join(notlar)


def calistir(dosyalar: list[Path], cikti: Path, platformlar: list[str] | None = None, hedef_roas: float | None = None,
             hedef_cpa: float | None = None, butce_yolu: Path | None = None, esleme: dict[str, str] | None = None) -> dict:
    tum, tanima, paralar = [], {}, set()
    for n, y in enumerate(dosyalar):
        p = platformlar[n] if platformlar and n < len(platformlar) else None
        satirlar, bulunan, para = dosya_oku(y, p, esleme or {})
        tum += satirlar
        tanima[y.name] = bulunan
        if para:
            paralar.add(para)
    uyarilar = []
    if len(paralar) > 1:
        uyarilar.append(f"Dosyalarda farklı para birimleri var ({', '.join(sorted(paralar))}); harcamalar toplanmadan önce çevrilmeli.")
    tarihli = [s["tarih"] for s in tum if s["tarih"]]
    toplam = olcutler({k: sum(s[k] for s in tum) for k in ("gosterim", "tiklama", "harcama", "donusum", "donusum_degeri")})
    platform = topla(tum, lambda s: s["platform"])
    kampanya = topla(tum, lambda s: (s["platform"], s["kampanya"]))
    for v in kampanya.values():
        v["not"] = isaretle(v, hedef_roas, hedef_cpa, toplam["harcama"])
    hafta = topla([s for s in tum if s["tarih"]], lambda s: (s["tarih"] - timedelta(days=s["tarih"].weekday())).isoformat())
    hafta = OrderedDict(sorted(hafta.items()))
    butce = butce_oku(butce_yolu, kampanya, tarihli) if butce_yolu else None
    _rapor(toplam, platform, kampanya, hafta, butce, tanima, uyarilar, hedef_roas, hedef_cpa, tarihli, paralar, cikti)
    return {"toplam": toplam, "platform": platform, "kampanya": kampanya, "hafta": hafta, "butce": butce, "uyarilar": uyarilar,
            "tanima": tanima}


def butce_oku(yol: Path, kampanya, tarihli) -> list[dict]:
    """Bütçe: Kampanya, Dönem Bütçesi, Başlangıç, Bitiş. Harcama temposu = harcanan ÷ (bütçe × geçen gün ÷ toplam gün)."""
    s = tablo_oku(yol)
    b = [katla(x) for x in s[0]]
    bul = lambda *a: next((b.index(katla(x)) for x in a if katla(x) in b), None)  # noqa: E731
    i_k, i_b = bul("kampanya", "campaign"), bul("bütçe", "dönem bütçesi", "budget")
    i_s, i_e = bul("başlangıç", "start"), bul("bitiş", "end")
    son_veri = max(tarihli) if tarihli else date.today()
    sonuc = []
    harcama = defaultdict(float)
    for (_, k), v in kampanya.items():
        harcama[katla(k)] += v["harcama"]
    for r in s[1:]:
        if not r[i_k]:
            continue
        bas, bit = tarih(r[i_s]) if i_s is not None else None, tarih(r[i_e]) if i_e is not None else None
        top_gun = (bit - bas).days + 1 if bas and bit else None
        gecen = min(top_gun, (son_veri - bas).days + 1) if top_gun else None
        bt = sayi(r[i_b])
        h = harcama.get(katla(r[i_k]), 0.0)
        beklenen = bt * gecen / top_gun if top_gun and gecen and gecen > 0 else None
        sonuc.append({"kampanya": str(r[i_k]), "butce": bt, "harcanan": h, "kalan": bt - h, "beklenen": beklenen,
                      "tempo": h / beklenen if beklenen else None})
    return sonuc


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")
OLCU_BAS = ["Harcama", "Gösterim", "Tıklama", "Dönüşüm", "Dönüşüm Değeri", "CTR", "TBM (CPC)", "BGBM (CPM)", "Dönüşüm Oranı", "CPA", "ROAS"]
OLCU_ALAN = ["harcama", "gosterim", "tiklama", "donusum", "donusum_degeri", "ctr", "cpc", "cpm", "cvr", "cpa", "roas"]
OLCU_BICIM = ["#,##0.00", "#,##0", "#,##0", "#,##0.0", "#,##0.00", "0.00%", "#,##0.00", "#,##0.00", "0.00%", "#,##0.00", "0.00"]


def _baslik(ws, satir=1):
    for c in ws[satir]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI


def _yaz(ws, anahtar: list, v: dict):
    ws.append(anahtar + [v[a] for a in OLCU_ALAN])
    for j, bc in enumerate(OLCU_BICIM, len(anahtar) + 1):
        ws.cell(ws.max_row, j).number_format = bc


def _rapor(toplam, platform, kampanya, hafta, butce, tanima, uyarilar, hedef_roas, hedef_cpa, tarihli, paralar, cikti):
    wb = Workbook()
    o = wb.active
    o.title = "Özet"
    donem = f"{min(tarihli):%d.%m.%Y} – {max(tarihli):%d.%m.%Y}" if tarihli else "tarih sütunu yok"
    o.append([f"Reklam kampanyası performansı · {donem}" + (f" · para birimi {', '.join(sorted(paralar))}" if paralar else "")])
    o["A1"].font = Font(bold=True, size=12)
    o.append(["Platform"] + OLCU_BAS)
    _baslik(o, 2)
    for k, v in platform.items():
        _yaz(o, [k], v)
    _yaz(o, ["TOPLAM"], toplam)
    for c in o[o.max_row]:
        c.font = Font(bold=True)
    o.append([])
    o.append(["Hedef ROAS", hedef_roas, "Hedef CPA", hedef_cpa])
    for u in uyarilar:
        o.append(["Uyarı", u])
    o.column_dimensions["A"].width = 18
    for j in range(2, len(OLCU_BAS) + 2):
        o.column_dimensions[get_column_letter(j)].width = 13

    k = wb.create_sheet("Kampanyalar")
    k.append(["Platform", "Kampanya"] + OLCU_BAS + ["Harcama Payı", "Değerlendirme"])
    _baslik(k)
    for (p, ad), v in kampanya.items():
        _yaz(k, [p, ad], v)
        k.cell(k.max_row, 14).value = v["harcama"] / toplam["harcama"] if toplam["harcama"] else None
        k.cell(k.max_row, 14).number_format = "0.0%"
        k.cell(k.max_row, 15).value = v["not"]
        if v["not"]:
            k.cell(k.max_row, 15).fill = YESIL if "artırılabilir" in v["not"] else SARI
    for j, w in enumerate([10, 34] + [12] * 11 + [10, 50], 1):
        k.column_dimensions[get_column_letter(j)].width = w
    k.freeze_panes = "C2"
    k.auto_filter.ref = k.dimensions
    n = min(15, len(kampanya))
    if n:
        g = BarChart()
        g.type, g.title, g.height, g.width = "bar", "Kampanya harcaması (ilk 15)", 9, 18
        g.add_data(Reference(k, min_col=3, min_row=1, max_row=1 + n), titles_from_data=True)
        g.set_categories(Reference(k, min_col=2, min_row=2, max_row=1 + n))
        o.add_chart(g, f"A{o.max_row + 3}")

    h = wb.create_sheet("Haftalık")
    h.append(["Hafta (Pazartesi)"] + OLCU_BAS)
    _baslik(h)
    for w, v in hafta.items():
        _yaz(h, [w], v)
    h.column_dimensions["A"].width = 16

    if butce:
        bt = wb.create_sheet("Bütçe Temposu")
        bt.append(["Kampanya", "Dönem Bütçesi", "Harcanan", "Kalan", "Bugüne Beklenen", "Tempo (harcanan / beklenen)"])
        _baslik(bt)
        for x in butce:
            bt.append([x["kampanya"], x["butce"], x["harcanan"], x["kalan"], x["beklenen"], x["tempo"]])
            for j in (2, 3, 4, 5):
                bt.cell(bt.max_row, j).number_format = "#,##0.00"
            bt.cell(bt.max_row, 6).number_format = "0%"
            if x["tempo"] is not None and not 0.85 <= x["tempo"] <= 1.15:
                bt.cell(bt.max_row, 6).fill = SARI
        bt.column_dimensions["A"].width = 34
        for c in "BCDEF":
            bt.column_dimensions[c].width = 16

    b = wb.create_sheet("Bilgi")
    for s in [["CTR", "tıklama ÷ gösterim"], ["TBM (CPC)", "harcama ÷ tıklama"], ["BGBM (CPM)", "harcama ÷ gösterim × 1000"],
              ["Dönüşüm oranı", "dönüşüm ÷ tıklama"], ["CPA", "harcama ÷ dönüşüm"], ["ROAS", "dönüşüm değeri ÷ harcama"],
              ["Dikkat", "Platformların dönüşüm tanımları ve ilişkilendirme (attribution) pencereleri farklıdır; aynı satış birden çok "
                         "platformda sayılabilir. Toplam ROAS'ı kendi sipariş verinizle (UTM) doğrulayın."]]:
        b.append(s)
    b.append([])
    b.append(["Dosya", "Tanınan sütunlar"])
    for d, t in tanima.items():
        b.append([d, " · ".join(f"{k}: {v}" for k, v in t.items() if v)])
    b.column_dimensions["A"].width = 22
    b.column_dimensions["B"].width = 130
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ornek = BURASI / "ornek_veri"
    ap = argparse.ArgumentParser(description="Reklam platformu dışa aktarımlarını birleştirip kampanya bazında performans raporu üretir.")
    ap.add_argument("--dosya", type=Path, action="append", help="Platform dışa aktarımı (.csv/.xlsx); birden çok verilebilir")
    ap.add_argument("--platform", action="append", help="Dosya sırasıyla platform adı (verilmezse dosya adı ve başlıklardan tahmin edilir)")
    ap.add_argument("--hedef-roas", type=float, help="Hedef ROAS (ör. 4)")
    ap.add_argument("--hedef-cpa", type=float, help="Hedef dönüşüm başı maliyet")
    ap.add_argument("--butce", type=Path, help="Bütçeler: Kampanya, Bütçe, Başlangıç, Bitiş")
    ap.add_argument("--esleme", nargs="*", help="Tanınmayan sütunlar için alan=başlık, ör. harcama='Spend (TRY)' donusum='Purchases'")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "reklam_performansi.xlsx")
    a = ap.parse_args(argv)
    if not a.dosya:
        a.dosya = [ornek / "google_ads_kampanya.csv", ornek / "meta_ads_kampanya.csv"]
        a.hedef_roas = a.hedef_roas or 4.0
        a.hedef_cpa = a.hedef_cpa or 150.0
        a.butce = a.butce or ornek / "butceler.csv"
    esleme = {}
    for e in a.esleme or []:
        k, _, v = e.partition("=")
        esleme[k.strip()] = v.strip().strip("'\"")
    s = calistir(a.dosya, a.cikti, a.platform, a.hedef_roas, a.hedef_cpa, a.butce, esleme)
    t = s["toplam"]
    print(f"[OK] {len(s['kampanya'])} kampanya · harcama {t['harcama']:,.2f} · dönüşüm {t['donusum']:,.0f} · "
          f"ROAS {(t['roas'] or 0):.2f} · CPA {(t['cpa'] or 0):,.2f}".replace(",", "X").replace(".", ",").replace("X", "."))
    for (p, k), v in s["kampanya"].items():
        if v["not"]:
            print(f"    [{p}] {k}: {v['not']}")
    for u in s["uyarilar"]:
        print(f"[!] {u}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
