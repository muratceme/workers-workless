"""
Masraf Fişi Kategorileme — Workers / Workless kod bloğu
Muhasebe › Muhasebe Elemanı

Çalışan masraf fişlerini ve kart harcamalarını kural dosyasına göre sınıflandırır, muhasebe kaydı önerir:
  - Kural: anahtar kelimeler (açıklama / satıcı) → kategori, alt hesap, KDV indirilebilir mi, KKEG oranı, limit.
    Öncelik numarası küçük olan kural önce denenir; ilk eşleşen kazanır. Kurallar kullanıcıya aittir.
  - Gider hesabı = departmanın fonksiyon hesabı (ör. 760 Pazarlama Satış Dağıtım, 770 Genel Yönetim) + alt hesap.
  - KDV: yalnız fatura türü belgelerde (e-Fatura, e-Arşiv, fatura, e-Bilet, serbest meslek makbuzu) ve kural izin
    veriyorsa 191 İndirilecek KDV'ye ayrılır; ÖKC fişinde indirim koşullara bağlı olduğundan gidere eklenir ve
    işaretlenir; belgesiz harcamada gider kabulü için belge gerekir (KKEG olarak işaretlenir).
  - Ödeme hesabı: Şirket Kartı → 309, Nakit (iş avansı) → 195, Kişisel Kart → 335 (değiştirilebilir).
  - Kontroller: sınıflanamayan fiş, KDV tutarı ile oranın tutarsızlığı, politika limiti aşımı, olası mükerrer
    (aynı çalışan + satıcı + tarih + tutar), hafta sonu harcaması, tanımsız departman.
Rapor: sınıflandırma, yevmiye önerisi, kategori / departman özeti, KKEG listesi, uyarılar. İnternete bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: 18 fiş, 12 kural
    python main.py --masraflar masraflar.xlsx --kurallar kurallar.csv --departmanlar departmanlar.csv
    python main.py --masraflar m.xlsx --kurallar k.csv --odeme-hesaplari "Şirket Kartı=300.05" "Nakit=195"
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
K2 = Decimal("0.01")
SIFIR = Decimal(0)
FATURA_TURLERI = ("fatura", "e fatura", "e arsiv", "e bilet", "bilet", "serbest meslek makbuzu", "e smm")
OKC_TURLERI = ("okc", "yazar kasa", "fis", "perakende")
BELGESIZ = ("belgesiz", "belge yok", "")
ODEME_VARSAYILAN = {"sirket karti": "309", "kurumsal kart": "309", "nakit": "195", "is avansi": "195", "kisisel kart": "335", "kisisel": "335",
                    "calisan": "335"}

MASRAF_SUTUNLARI = {"no": ("fis no", "masraf no", "belge no"), "tarih": ("tarih", "harcama tarihi"), "calisan": ("calisan", "ad soyad", "personel"),
                    "departman": ("departman", "birim", "masraf merkezi"), "satici": ("satici", "firma", "isyeri", "satici unvan"),
                    "aciklama": ("aciklama", "harcama aciklamasi"), "tutar": ("tutar kdv dahil", "tutar", "toplam", "kdv dahil tutar"),
                    "oran": ("kdv orani",), "kdv": ("kdv tutari", "kdv"), "belge": ("belge turu", "belge"), "odeme": ("odeme", "odeme sekli", "odeme turu")}
KURAL_SUTUNLARI = {"oncelik": ("oncelik", "sira"), "kelimeler": ("anahtar kelimeler", "anahtar kelime", "kelimeler"), "alan": ("alan",),
                   "kategori": ("kategori",), "alt": ("alt hesap", "hesap", "hesap kodu"), "kdv": ("kdv indirilebilir",), "kkeg": ("kkeg", "kkeg orani"),
                   "limit": ("limit tl", "limit"), "not": ("not", "aciklama")}
DEP_SUTUNLARI = {"departman": ("departman", "birim"), "hesap": ("gider hesabi", "hesap", "fonksiyon hesabi")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("%", "").replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tl(x: Decimal | None) -> str:
    return "—" if x is None else f"{x.quantize(K2, ROUND_HALF_UP):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Kural:
    oncelik: int
    kelimeler: list
    alan: str                  # "aciklama" | "satici" | "hepsi"
    kategori: str
    alt: str
    kdv: bool
    kkeg: Decimal
    limit: Decimal | None
    not_: str

    def eslesir(self, aciklama: str, satici: str) -> str | None:
        hedef = {"aciklama": aciklama, "satici": satici}.get(self.alan, f"{aciklama} {satici}")
        h = f" {katla(hedef)} "
        return next((k for k in self.kelimeler if f" {katla(k)} " in h or (len(katla(k)) >= 5 and katla(k) in h)), None)


@dataclass
class Masraf:
    no: str
    tarih: date | None
    calisan: str
    departman: str
    satici: str
    aciklama: str
    tutar: Decimal
    oran: Decimal | None
    kdv: Decimal
    belge: str
    odeme: str
    kural: Kural | None = None
    kelime: str = ""
    gider_hesabi: str = ""
    odeme_hesabi: str = ""
    indirilecek_kdv: Decimal = SIFIR
    kkeg: Decimal = SIFIR
    bulgular: list = field(default_factory=list)       # (önem, tür, açıklama)

    @property
    def gider(self) -> Decimal:
        return self.tutar - self.indirilecek_kdv

    @property
    def belge_turu(self) -> str:
        b = katla(self.belge)
        if any(b == x or b.startswith(x) or x in b for x in FATURA_TURLERI if x):
            return "fatura"
        if any(x in b for x in OKC_TURLERI):
            return "okc"
        if b in BELGESIZ:
            return "belgesiz"
        return "diger"


def oku_kurallar(yol: Path) -> list[Kural]:
    k = []
    for r in kayitlar(yol, KURAL_SUTUNLARI, ("kelimeler", "kategori")):
        kel = [x.strip() for x in metin(r.get("kelimeler")).split("|") if x.strip()]
        if not kel:
            continue
        alan = katla(r.get("alan"))
        k.append(Kural(int(para(r.get("oncelik")) or 999), kel, "aciklama" if alan.startswith("acik") else "satici" if alan.startswith("satic") else "hepsi",
                       metin(r["kategori"]), metin(r.get("alt")), not katla(r.get("kdv")).startswith("h"), para(r.get("kkeg")) or SIFIR, para(r.get("limit")),
                       metin(r.get("not"))))
    return sorted(k, key=lambda x: x.oncelik)


def oku_masraflar(yol: Path) -> tuple[list[Masraf], list[dict]]:
    m, uy = [], []
    for r in kayitlar(yol, MASRAF_SUTUNLARI, ("tutar",)):
        t = para(r.get("tutar"))
        if t is None:
            if any(r.get(k) for k in ("no", "aciklama")):
                uy.append({"onem": "Orta", "tur": "Okunamayan satır", "fis": metin(r.get("no")), "aciklama": f"Satır {r['_satir']}: tutar okunamadı"})
            continue
        oran = para(r.get("oran"))
        m.append(Masraf(metin(r.get("no")) or f"S{r['_satir']}", tarih(r.get("tarih")), metin(r.get("calisan")), metin(r.get("departman")), metin(r.get("satici")),
                        metin(r.get("aciklama")), t, oran, para(r.get("kdv")) or SIFIR, metin(r.get("belge")), metin(r.get("odeme"))))
    return m, uy


# ----------------------------------------------------------------------------
# Sınıflandırma
# ----------------------------------------------------------------------------

def siniflandir(masraflar: list[Masraf], kurallar: list[Kural], departmanlar: dict[str, str], odeme_hesaplari: dict[str, str],
                varsayilan_hesap: str = "770") -> None:
    for m in masraflar:
        b = m.bulgular
        for k in kurallar:
            kel = k.eslesir(m.aciklama, m.satici)
            if kel:
                m.kural, m.kelime = k, kel
                break
        fonk = departmanlar.get(katla(m.departman))
        if fonk is None:
            fonk = varsayilan_hesap
            b.append(("Orta", "Tanımsız departman", f"'{m.departman or '—'}' departman listesinde yok; {varsayilan_hesap} kullanıldı"))
        if m.kural is None:
            m.gider_hesabi = f"{fonk}.??"
            b.append(("Yüksek", "Sınıflanamadı", "Hiçbir kurala uymadı; kategoriyi elle seçin ve gerekirse kural ekleyin"))
        else:
            m.gider_hesabi = f"{fonk}.{m.kural.alt}" if m.kural.alt and "." not in m.kural.alt and len(m.kural.alt) < 3 else (m.kural.alt or fonk)
        m.odeme_hesabi = odeme_hesaplari.get(katla(m.odeme), "")
        if not m.odeme_hesabi:
            b.append(("Orta", "Ödeme şekli", f"'{m.odeme or '—'}' için ödeme hesabı tanımlı değil (--odeme-hesaplari)"))
        # KDV
        if m.oran is not None and m.oran > 0:
            beklenen = (m.tutar * m.oran / (100 + m.oran)).quantize(K2, ROUND_HALF_UP)
            if abs(beklenen - m.kdv) > Decimal("0.05"):
                b.append(("Orta", "KDV tutarsız", f"{tl(m.tutar)} × {m.oran:g}/{100 + m.oran:g} = {tl(beklenen)}; fişte {tl(m.kdv)}"))
        tur = m.belge_turu
        if m.kdv > 0:
            if m.kural and not m.kural.kdv:
                b.append(("Bilgi", "KDV indirilmez", f"'{m.kural.kategori}' kuralında KDV indirilemez; gidere eklendi"))
            elif tur == "fatura":
                m.indirilecek_kdv = m.kdv
            elif tur == "okc":
                b.append(("Bilgi", "ÖKC fişi KDV", "ÖKC fişindeki KDV'nin indirimi koşullara bağlıdır; gidere eklendi, mali müşavirle doğrulayın"))
            else:
                b.append(("Bilgi", "KDV indirilmez", f"Belge türü '{m.belge or '—'}' KDV indirimine uygun değil; gidere eklendi"))
        if tur == "belgesiz":
            m.kkeg = m.gider
            b.append(("Yüksek", "Belgesiz harcama", "Gider yazmak için belge gerekir; belge alınamazsa kanunen kabul edilmeyen gider (KKEG) sayılır"))
        elif m.kural and m.kural.kkeg:
            m.kkeg = (m.gider * m.kural.kkeg / 100).quantize(K2, ROUND_HALF_UP)
            b.append(("Bilgi", "KKEG", f"{m.kural.kategori}: giderin %{m.kural.kkeg:g}'i KKEG ({tl(m.kkeg)} TL). {m.kural.not_}".strip()))
        if m.kural and m.kural.limit and m.tutar > m.kural.limit:
            b.append(("Orta", "Limit aşımı", f"{m.kural.kategori} politika limiti {tl(m.kural.limit)} TL; fiş {tl(m.tutar)} TL "
                      "(gece / kişi sayısına bölerek kontrol edin)"))
        if m.kural and m.kural.not_ and not m.kural.kkeg:
            b.append(("Bilgi", "Kural notu", m.kural.not_))
        if m.tarih and m.tarih.weekday() >= 5:
            b.append(("Bilgi", "Hafta sonu", f"{m.tarih:%d.%m.%Y} hafta sonu; iş amacını kontrol edin"))
    anahtar = Counter((katla(m.calisan), katla(m.satici), m.tarih, m.tutar) for m in masraflar)
    for m in masraflar:
        if anahtar[(katla(m.calisan), katla(m.satici), m.tarih, m.tutar)] > 1:
            m.bulgular.append(("Yüksek", "Olası mükerrer", "Aynı çalışan, satıcı, tarih ve tutarla birden fazla fiş var"))


def yevmiye(masraflar: list[Masraf]) -> list[list]:
    satirlar = []
    for m in masraflar:
        satirlar.append([m.tarih, m.no, m.gider_hesabi, (m.kural.kategori if m.kural else "Sınıflanamadı") + f" · {m.aciklama}", m.gider, SIFIR])
        if m.indirilecek_kdv:
            satirlar.append([m.tarih, m.no, "191", "İndirilecek KDV", m.indirilecek_kdv, SIFIR])
        satirlar.append([m.tarih, m.no, m.odeme_hesabi or "???", f"{m.odeme} · {m.calisan}", SIFIR, m.tutar])
    return satirlar


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KONTROL = PatternFill("solid", fgColor="FFF4CE")
UST = Alignment(vertical="top", wrap_text=True)
PF = "#,##0.00"
SIRA = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def rapor_yaz(cikti: Path, s: dict) -> None:
    ms = s["masraflar"]
    wb = Workbook()
    sn = wb.active
    sn.title = "Sınıflandırma"
    _baslik(sn, ["Fiş No", "Tarih", "Çalışan", "Departman", "Satıcı", "Açıklama", "Tutar", "Belge", "Kategori", "Eşleşen Kelime", "Gider Hesabı",
                 "Gider", "İndirilecek KDV", "KKEG", "Ödeme Hesabı", "Bulgular", "Onay / Düzeltme"],
            (9, 11, 13, 12, 22, 34, 11, 13, 18, 14, 11, 11, 11, 10, 9, 44, 16))
    for m in ms:
        sn.append([m.no, m.tarih, m.calisan, m.departman, m.satici, m.aciklama, float(m.tutar), m.belge, m.kural.kategori if m.kural else "—", m.kelime,
                   m.gider_hesabi, float(m.gider), float(m.indirilecek_kdv), float(m.kkeg), m.odeme_hesabi,
                   "\n".join(f"{t}: {a}" for o, t, a in sorted(m.bulgular, key=lambda x: SIRA[x[0]]) if o != "Bilgi" or t in ("KKEG", "ÖKC fişi KDV")), ""])
        r = sn.max_row
        sn.cell(r, 2).number_format = "DD.MM.YYYY"
        for j in (7, 12, 13, 14):
            sn.cell(r, j).number_format = PF
        en = min((b[0] for b in m.bulgular), key=SIRA.get, default=None)
        if en in ("Yüksek", "Orta"):
            sn.cell(r, 16).fill = PatternFill("solid", fgColor=RENK[en])
        sn.cell(r, 16).alignment = UST
        sn.cell(r, 17).fill = KONTROL
    sn.auto_filter.ref = f"A1:Q{sn.max_row}"

    yv = wb.create_sheet("Yevmiye Önerisi")
    _baslik(yv, ["Tarih", "Fiş No", "Hesap", "Açıklama", "Borç", "Alacak"], (11, 9, 11, 60, 13, 13))
    for t, no, h, a, b, al in s["yevmiye"]:
        yv.append([t, no, h, a, float(b) or None, float(al) or None])
        yv.cell(yv.max_row, 1).number_format = "DD.MM.YYYY"
        for j in (5, 6):
            yv.cell(yv.max_row, j).number_format = PF
    tb = sum((x[4] for x in s["yevmiye"]), SIFIR)
    ta = sum((x[5] for x in s["yevmiye"]), SIFIR)
    yv.append(["", "", "", "Toplam", float(tb), float(ta)])
    for j in (5, 6):
        yv.cell(yv.max_row, j).number_format = PF
        yv.cell(yv.max_row, j).font = Font(bold=True)

    oz = wb.create_sheet("Özet")
    _baslik(oz, ["Kategori", "Fiş", "Tutar", "Gider", "İndirilecek KDV", "KKEG"], (24, 6, 14, 14, 14, 12))
    g = defaultdict(list)
    for m in ms:
        g[m.kural.kategori if m.kural else "Sınıflanamadı"].append(m)
    for ad, lst in sorted(g.items(), key=lambda i: -sum(m.tutar for m in i[1])):
        oz.append([ad, len(lst), float(sum(m.tutar for m in lst)), float(sum(m.gider for m in lst)), float(sum(m.indirilecek_kdv for m in lst)),
                   float(sum(m.kkeg for m in lst))])
        for j in (3, 4, 5, 6):
            oz.cell(oz.max_row, j).number_format = PF
    oz.append([])
    _baslik(oz, ["Departman", "Fiş", "Tutar", "Gider", "İndirilecek KDV", "KKEG"], ())
    g = defaultdict(list)
    for m in ms:
        g[m.departman or "—"].append(m)
    for ad, lst in sorted(g.items()):
        oz.append([ad, len(lst), float(sum(m.tutar for m in lst)), float(sum(m.gider for m in lst)), float(sum(m.indirilecek_kdv for m in lst)),
                   float(sum(m.kkeg for m in lst))])
        for j in (3, 4, 5, 6):
            oz.cell(oz.max_row, j).number_format = PF

    kk = wb.create_sheet("KKEG Listesi")
    _baslik(kk, ["Fiş No", "Tarih", "Çalışan", "Kategori", "Açıklama", "Gider", "KKEG", "Gerekçe"], (9, 11, 13, 18, 34, 12, 12, 60))
    for m in ms:
        if m.kkeg:
            ger = next((a for o, t, a in m.bulgular if t in ("KKEG", "Belgesiz harcama")), "")
            kk.append([m.no, m.tarih, m.calisan, m.kural.kategori if m.kural else "—", m.aciklama, float(m.gider), float(m.kkeg), ger])
            kk.cell(kk.max_row, 2).number_format = "DD.MM.YYYY"
            for j in (6, 7):
                kk.cell(kk.max_row, j).number_format = PF
            kk.cell(kk.max_row, 8).alignment = UST

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Fiş No", "Açıklama"], (9, 20, 9, 100))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["fis"], u["aciklama"]])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(masraf_yolu: Path, kural_yolu: Path, cikti: Path, dep_yolu: Path | None = None, odeme: dict[str, str] | None = None,
             varsayilan_hesap: str = "770") -> dict:
    masraflar, uyarilar = oku_masraflar(masraf_yolu)
    if not masraflar:
        raise ValueError(f"{masraf_yolu.name}: masraf bulunamadı")
    kurallar = oku_kurallar(kural_yolu)
    deps = {katla(r.get("departman")): metin(r.get("hesap")) for r in (kayitlar(dep_yolu, DEP_SUTUNLARI, ("departman", "hesap")) if dep_yolu else [])
            if metin(r.get("departman"))}
    oh = dict(ODEME_VARSAYILAN)
    oh.update({katla(k): v for k, v in (odeme or {}).items()})
    siniflandir(masraflar, kurallar, deps, oh, varsayilan_hesap)
    uyarilar += [{"onem": o, "tur": t, "fis": m.no, "aciklama": a} for m in masraflar for o, t, a in m.bulgular if t != "Kural notu"]
    uyarilar.sort(key=lambda u: (SIRA[u["onem"]], u["tur"], u["fis"]))
    s = {"masraflar": masraflar, "kurallar": kurallar, "uyarilar": uyarilar, "yevmiye": yevmiye(masraflar)}
    rapor_yaz(cikti, s)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Masraf fişlerini kurallara göre sınıflandırır ve muhasebe kaydı önerir.")
    p.add_argument("--masraflar", type=Path, default=ORNEK / "masraflar.csv",
                   help="Fiş No, Tarih, Çalışan, Departman, Satıcı, Açıklama, Tutar (KDV dahil), KDV Oranı, KDV Tutarı, Belge Türü, Ödeme")
    p.add_argument("--kurallar", type=Path, default=ORNEK / "kurallar.csv",
                   help="Öncelik, Anahtar Kelimeler (| ile), Alan, Kategori, Alt Hesap, KDV İndirilebilir, KKEG (%%), Limit (TL), Not")
    p.add_argument("--departmanlar", type=Path, help="Departman, Gider Hesabı (ör. Satış;760)")
    p.add_argument("--varsayilan-hesap", default="770", help="Departmanı tanımsız fişlerde fonksiyon hesabı (varsayılan 770)")
    p.add_argument("--odeme-hesaplari", nargs="*", default=[], metavar="ÖDEME=HESAP", help="ör. \"Şirket Kartı=309\" \"Nakit=195\" \"Kişisel Kart=335\"")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "masraf_kategorileme.xlsx")
    a = p.parse_args(argv)
    dep = a.departmanlar or (ORNEK / "departmanlar.csv" if a.masraflar == ORNEK / "masraflar.csv" else None)
    for y in (a.masraflar, a.kurallar, dep):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    odeme = {}
    for x in a.odeme_hesaplari:
        k, _, v = x.partition("=")
        if not v.strip():
            print(f"[X] --odeme-hesaplari 'Ödeme=Hesap' biçiminde olmalı: {x}")
            return 2
        odeme[k.strip()] = v.strip()
    try:
        s = calistir(a.masraflar, a.kurallar, a.cikti, dep, odeme, a.varsayilan_hesap)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    ms = s["masraflar"]
    print(f"[OK] {len(ms)} fiş · {sum(1 for m in ms if m.kural)} sınıflandı · toplam {tl(sum((m.tutar for m in ms), SIFIR))} TL · "
          f"indirilecek KDV {tl(sum((m.indirilecek_kdv for m in ms), SIFIR))} TL · KKEG {tl(sum((m.kkeg for m in ms), SIFIR))} TL")
    for u in s["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['fis']} {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
