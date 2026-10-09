import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "m.xlsx"
        cls.s = main.calistir(ORNEK / "masraflar.csv", ORNEK / "kurallar.csv", cls.cikti, ORNEK / "departmanlar.csv")
        cls.m = {m.no: m for m in cls.s["masraflar"]}
        cls.u = {(u["fis"], u["tur"]) for u in cls.s["uyarilar"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_siniflandirma(self):
        k = {no: (m.kural.kategori if m.kural else None, m.gider_hesabi) for no, m in self.m.items()}
        self.assertEqual(k["M-1003"], ("Temsil ve ağırlama", "760.06"))      # "müşteri yemeği" kuralı "yemek"ten önce
        self.assertEqual(k["M-1005"], ("Binek araç gideri", "760.05"))
        self.assertEqual(k["M-1006"], ("Yakıt (ticari araç)", "730.04"))
        self.assertEqual(k["M-1011"], ("Yazılım aboneliği", "750.11"))
        self.assertEqual(k["M-1013"], (None, "770.??"))
        self.assertEqual(sum(1 for m in self.m.values() if m.kural), 17)

    def test_kdv_ve_kkeg(self):
        m = self.m
        self.assertEqual((m["M-1001"].indirilecek_kdv, m["M-1001"].gider), (D("600.00"), D("6000.00")))
        self.assertEqual(m["M-1002"].indirilecek_kdv, 0)                     # ÖKC fişi → gidere
        self.assertEqual(m["M-1010"].indirilecek_kdv, 0)
        self.assertEqual(m["M-1005"].kkeg, D("700.00"))                       # 2.333,33 × %30
        self.assertEqual((m["M-1007"].kkeg, m["M-1012"].kkeg), (D("2167.00"), D("850.00")))
        self.assertEqual((m["M-1006"].odeme_hesabi, m["M-1007"].odeme_hesabi, m["M-1001"].odeme_hesabi), ("195", "335", "309"))

    def test_uyarilar(self):
        self.assertTrue({("M-1012", "Belgesiz harcama"), ("M-1014", "Olası mükerrer"), ("M-1015", "Olası mükerrer"), ("M-1013", "Sınıflanamadı"),
                         ("M-1016", "KDV tutarsız"), ("M-1018", "Tanımsız departman"), ("M-1001", "Limit aşımı"), ("M-1002", "Limit aşımı")} <= self.u)
        self.assertNotIn(("M-1003", "Limit aşımı"), self.u)

    def test_yevmiye_ve_excel(self):
        yv = self.s["yevmiye"]
        self.assertEqual(sum(x[4] for x in yv), sum(x[5] for x in yv))
        self.assertEqual(sum(1 for x in yv if x[2] == "191"), 11)
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Sınıflandırma", "Yevmiye Önerisi", "Özet", "KKEG Listesi", "Uyarılar"])
        self.assertEqual(wb["KKEG Listesi"].max_row, 4)


class KuralTesti(unittest.TestCase):
    def test_eslesme_ve_odeme(self):
        k = main.Kural(1, ["bp", "akaryakıt"], "hepsi", "Yakıt", "04", True, D(0), None, "")
        self.assertEqual(k.eslesir("BP istasyon", ""), "bp")
        self.assertIsNone(k.eslesir("bpm yazılım", ""))
        self.assertEqual(k.eslesir("Akaryakıt alımı", ""), "akaryakıt")
        with tempfile.TemporaryDirectory() as tmp:
            s = main.calistir(ORNEK / "masraflar.csv", ORNEK / "kurallar.csv", Path(tmp) / "o.xlsx", None, {"Şirket Kartı": "300.05"}, "740")
            m = {x.no: x for x in s["masraflar"]}
            self.assertEqual((m["M-1001"].odeme_hesabi, m["M-1001"].gider_hesabi), ("300.05", "740.02"))

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--odeme-hesaplari", "Nakit"]), 2)
            self.assertEqual(main.main(["--kurallar", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
