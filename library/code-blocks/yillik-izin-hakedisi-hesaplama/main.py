"""
Yıllık İzin Hakedişi Hesaplama — Workers / Workless kod bloğu
İnsan Kaynakları › Bordro ve Özlük İşleri Uzmanı

4857 sayılı İş Kanunu md. 53 kurallarıyla her çalışanın yıllık ücretli izin hakkını hizmet yılı
bazında hesaplar; toplam hakediş, kullanılan, kalan izin, sonraki hakediş tarihi ve (brüt verilirse)
kullanılmayan izin ücreti tahminini Excel'e yazar. İnternete bağlanmaz.

Kurallar (md. 53):
  - Hizmet süresi 1-5 yıl (5 dahil): 14 gün · 5 yıldan fazla 15 yıldan az: 20 gün · 15 yıl ve üzeri: 26 gün
  - 18 yaş ve altı ile 50 yaş ve üstü işçilere en az 20 gün
  - Yer altı işlerinde çalışanlara +4 gün
  - Hak, her hizmet yılı tamamlandığında doğar (deneme süresi dahil)

Kullanım:
    python main.py                                         # örnek personel listesiyle dener
    python main.py --girdi personel.xlsx --tarih 31.12.2026
"""
from __future__ import annotations

import argparse
import csv
import sys
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent


def tarih(x) -> date | None:
    if x in (None, ""):
        return None
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x).strip()
    for f in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    raise ValueError(f"Tarih anlaşılamadı: {x!r}")


def yil_ekle(t: date, n: int) -> date:
    """n yıl sonrası; 29 Şubat artık olmayan yıla 28 Şubat olarak taşınır."""
    try:
        return t.replace(year=t.year + n)
    except ValueError:
        return t.replace(year=t.year + n, day=28)


def yas(dogum: date, gun: date) -> int:
    return gun.year - dogum.year - ((gun.month, gun.day) < (dogum.month, dogum.day))


def yillik_izin_gunu(hizmet_yili: int, yas_: int | None, yer_alti: bool = False) -> int:
    """hizmet_yili: o hakedişte tamamlanan hizmet yılı (1, 2, 3...)."""
    if hizmet_yili <= 5:
        gun = 14
    elif hizmet_yili < 15:
        gun = 20
    else:
        gun = 26
    if yas_ is not None and (yas_ <= 18 or yas_ >= 50):
        gun = max(gun, 20)
    return gun + (4 if yer_alti else 0)


@dataclass
class Hakedis:
    hizmet_yili: int
    tarih: date
    yas: int | None
    gun: int


@dataclass
class Calisan:
    ad: str
    giris: date
    dogum: date | None = None
    kullanilan: Decimal = Decimal(0)
    devreden: Decimal = Decimal(0)
    yer_alti: bool = False
    brut: Decimal | None = None
    hakedisler: list = field(default_factory=list)
    notlar: list = field(default_factory=list)


def hesapla(c: Calisan, referans: date) -> Calisan:
    if c.giris > referans:
        c.notlar.append("İşe giriş tarihi referans tarihinden sonra")
        return c
    n = 1
    while (t := yil_ekle(c.giris, n)) <= referans:
        y = yas(c.dogum, t) if c.dogum else None
        c.hakedisler.append(Hakedis(n, t, y, yillik_izin_gunu(n, y, c.yer_alti)))
        n += 1
    if not c.dogum:
        c.notlar.append("Doğum tarihi yok: 18 yaş altı / 50 yaş üstü kuralı uygulanamadı")
    if not c.hakedisler:
        c.notlar.append("Henüz 1 hizmet yılı dolmadı; izin hakkı doğmadı")
    return c


def ozet(c: Calisan, referans: date) -> dict:
    toplam = sum(h.gun for h in c.hakedisler) + c.devreden
    kalan = toplam - c.kullanilan
    sonraki_n = len(c.hakedisler) + 1
    sonraki_t = yil_ekle(c.giris, sonraki_n)
    sonraki_yas = yas(c.dogum, sonraki_t) if c.dogum else None
    d = {
        "Ad Soyad": c.ad, "İşe Giriş": c.giris, "Doğum Tarihi": c.dogum,
        "Tamamlanan Hizmet Yılı": len(c.hakedisler),
        "Son Hakediş Tarihi": c.hakedisler[-1].tarih if c.hakedisler else None,
        "Son Hakediş (Gün)": c.hakedisler[-1].gun if c.hakedisler else 0,
        "Devreden": float(c.devreden), "Toplam Hakediş": float(toplam), "Kullanılan": float(c.kullanilan),
        "Kalan": float(kalan),
        "Sonraki Hakediş Tarihi": sonraki_t,
        "Sonraki Hakediş (Gün)": yillik_izin_gunu(sonraki_n, sonraki_yas, c.yer_alti),
        "Kullanılmayan İzin Ücreti (Brüt)": None,
        "Notlar": "; ".join(c.notlar + (["Kullanılan izin hakedişi aşıyor (avans izin)"] if kalan < 0 else [])),
    }
    if c.brut is not None and kalan > 0:
        d["Kullanılmayan İzin Ücreti (Brüt)"] = float((c.brut / 30 * kalan).quantize(Decimal("0.01"), ROUND_HALF_UP))
    return d


# ----------------------------------------------------------------------------
# Girdi / çıktı
# ----------------------------------------------------------------------------

def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


SUTUNLAR = {
    "ad": ("ad soyad", "ad", "personel", "çalışan", "sicil"),
    "giris": ("işe giriş", "işe giriş tarihi", "giriş tarihi", "giriş", "kıdem başlangıç"),
    "dogum": ("doğum tarihi", "doğum"),
    "kullanilan": ("kullanılan izin", "kullanılan", "kullanılan gün"),
    "devreden": ("devreden izin", "devreden", "önceki yıllardan devreden"),
    "yer_alti": ("yer altı", "yeraltı"),
    "brut": ("brüt", "brüt ücret"),
}


def sayi(x) -> Decimal:
    if x in (None, ""):
        return Decimal(0)
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    return Decimal(s)


def liste_oku(yol: Path) -> list[Calisan]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ayirici = csv.Sniffer().sniff(metin.splitlines()[0], delimiters=",;\t").delimiter
        satirlar = list(csv.reader(metin.splitlines(), delimiter=ayirici))
    satirlar = [r for r in satirlar if any(c not in (None, "") for c in r)]
    b = [kucuk(x) for x in satirlar[0]]
    k = {alan: next((i for i, x in enumerate(b) if x in adlar), None) for alan, adlar in SUTUNLAR.items()}
    if k["ad"] is None or k["giris"] is None:
        raise SystemExit(f"Gerekli sütunlar: Ad Soyad ve İşe Giriş. Bulunan: {satirlar[0]}")
    al = lambda r, alan: r[k[alan]] if k[alan] is not None else None  # noqa: E731
    return [Calisan(
        ad=str(al(r, "ad")).strip(), giris=tarih(al(r, "giris")), dogum=tarih(al(r, "dogum")),
        kullanilan=sayi(al(r, "kullanilan")), devreden=sayi(al(r, "devreden")),
        yer_alti=kucuk(al(r, "yer_alti")) in {"e", "evet", "1", "true"},
        brut=sayi(al(r, "brut")) if al(r, "brut") not in (None, "") else None,
    ) for r in satirlar[1:]]


BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")


def rapor_yaz(calisanlar: list[Calisan], referans: date, cikti: Path) -> list[dict]:
    wb = Workbook()
    ws = wb.active
    ws.title = "İzin Durumu"
    ozetler = [ozet(c, referans) for c in calisanlar]
    basliklar = list(ozetler[0]) if ozetler else []
    ws.append(basliklar)
    for h in ws[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for o in ozetler:
        ws.append([o[b] for b in basliklar])
        if o["Kalan"] < 0:
            for h in ws[ws.max_row]:
                h.fill = KIRMIZI
    for j, b in enumerate(basliklar, 1):
        ws.column_dimensions[get_column_letter(j)].width = 50 if b == "Notlar" else max(12, len(b) + 2)
        for h in ws[get_column_letter(j)][1:]:
            if isinstance(h.value, date):
                h.number_format = "DD.MM.YYYY"
    ws.freeze_panes = "B2"

    d = wb.create_sheet("Hakediş Detayı")
    d.append(["Ad Soyad", "Hizmet Yılı", "Hakediş Tarihi", "O Tarihteki Yaş", "İzin Günü"])
    for h in d[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for c in calisanlar:
        for h in c.hakedisler:
            d.append([c.ad, h.hizmet_yili, h.tarih, h.yas, h.gun])
            d.cell(d.max_row, 3).number_format = "DD.MM.YYYY"
    for col, g in zip("ABCDE", (24, 12, 16, 16, 12)):
        d.column_dimensions[col].width = g

    k = wb.create_sheet("Kurallar")
    for s in [["Referans tarih", referans.strftime("%d.%m.%Y")],
              ["Dayanak", "4857 sayılı İş Kanunu md. 53"],
              ["1-5 yıl (5 dahil)", "14 gün"], ["5 yıldan fazla, 15 yıldan az", "20 gün"], ["15 yıl ve üzeri", "26 gün"],
              ["Yaş kuralı", "Hakediş tarihinde 18 yaş ve altı ile 50 yaş ve üstü olanlara en az 20 gün"],
              ["Yer altı işleri", "+4 gün"],
              ["Not", "Toplu iş sözleşmesi veya iş sözleşmesiyle daha uzun süre verilebilir; bu araç kanuni asgariyi hesaplar."]]:
        k.append(s)
    k.column_dimensions["A"].width = 30
    k.column_dimensions["B"].width = 90
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    return ozetler


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description="Yıllık ücretli izin hakedişi hesabı (4857 s. Kanun md. 53).")
    ap.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "personel.csv")
    ap.add_argument("--tarih", default=None, help="Referans tarih (varsayılan: bugün)")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "yillik_izin.xlsx")
    a = ap.parse_args(argv)
    referans = tarih(a.tarih) if a.tarih else date.today()
    calisanlar = [hesapla(c, referans) for c in liste_oku(a.girdi)]
    ozetler = rapor_yaz(calisanlar, referans, a.cikti)
    for o in ozetler:
        print(f"[OK] {o['Ad Soyad']}: hakediş {o['Toplam Hakediş']:g} · kullanılan {o['Kullanılan']:g} · kalan {o['Kalan']:g} gün")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
