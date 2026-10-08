import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))

import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"


class YaziTesti(unittest.TestCase):
    def test_yazidan_tutar(self):
        for yazi, beklenen in [("Yalnız Yüzyirmibeşbin Türk Lirası Elli Kuruş", "125000.50"), ("Bir milyon iki yüz bin TL", "1200000"),
                               ("Seksenikibinbeşyüz TL", "82500"), ("Üçbinyediyüzelli TL Yirmibeş Kuruş", "3750.25"),
                               ("Yüzyirmibeşbin TL Elli Kr.", "125000.50"), ("dokuzyüzdoksandokuzbindokuzyüzdoksandokuz", "999999"),
                               ("altmışaltı", "66"), ("yedimilyarbeşyüzmilyon", "7500000000"), ("Bin TL", "1000"), ("Onbin Dolar", "10000"),
                               ("#125.000,50#", "125000.50")]:
            self.assertEqual(main.yazidan_tutar(yazi), D(beklenen), yazi)

    def test_okunamayan(self):
        self.assertEqual(main.yazidan_tutar("Yüz yirmi beş bin lira kırk beş"), D("125000.45"))   # "kuruş" yazılmasa da
        self.assertIsNone(main.yazidan_tutar("Yüzbin TL elli kuruş on"))                           # kuruştan sonra sayı
        self.assertIsNone(main.yazidan_tutar("Yüzbin TL ellikuruş fazla"))
        self.assertIsNone(main.yazidan_tutar("Yüz TL yüzelli Kuruş"))

    def test_iban(self):
        self.assertEqual(main.iban_kontrol("TR480099900001000000100001"), "")
        self.assertIn("mod 97", main.iban_kontrol("TR480099900001000000100002"))
        self.assertIn("26 karakter", main.iban_kontrol("TR48009990000100000010000"))
        self.assertIn("rezerv", main.iban_kontrol("TR480099910001000000100001"))
        self.assertEqual(main.iban_kontrol("DE89370400440532013000"), "")             # ISO 13616 örnek IBAN
        self.assertEqual(main.iban_temizle("tr48 0099 9000 0100 0000 1000 01"), "TR480099900001000000100001")

    def test_unvan(self):
        self.assertEqual(main.unvan_karsilastir("Örnek Gıda A.Ş.", "ÖRNEK GIDA SAN. VE TİC. A.Ş."), "")
        self.assertEqual(main.unvan_karsilastir("Örnek Gıda", "Örnek Gıda Pazarlama A.Ş."), "kisa")
        self.assertEqual(main.unvan_karsilastir("Örnek Tarım", "Kurgu Lojistik"), "farkli")


class OrnekTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        with tempfile.TemporaryDirectory() as tmp:
            cls.s = main.calistir(ORNEK / "talimatlar.csv", Path(tmp) / "t.xlsx", ORNEK / "hesaplar.csv", ORNEK / "imza_yetkileri.csv")
            cls.wb = load_workbook(Path(tmp) / "t.xlsx")
        cls.t = {t.no: t for t in cls.s["talimatlar"]}

    def bul(self, no):
        return "\n".join(f"{o} {a}: {m}" for o, a, m in self.t[no].bulgular)

    def test_durumlar(self):
        self.assertEqual(self.s["banka"], "00999")
        durum = {n: t.durum for n, t in self.t.items()}
        self.assertEqual({n for n, d in durum.items() if d == "Uygun"}, {"T01", "T05", "T06", "T12", "T15"})
        self.assertEqual({n for n, d in durum.items() if d == "Kontrol"}, {"T11"})

    def test_bulgular(self):
        self.assertIn("Hata Alıcı IBAN: IBAN kontrol hanesi", self.bul("T02"))
        self.assertIn("münferit limiti (500.000,00 TL) tutarın (750.000,00 TL) altında", self.bul("T03"))
        self.assertIn("Murat Kurgu: yetki süresi 30.06.2026 tarihinde dolmuş", self.bul("T04"))
        self.assertIn("ikinci bir yetkilinin imzası gerekli", self.bul("T04"))
        self.assertIn("Virman yalnız bankamızdaki hesaplar arasında", self.bul("T07"))
        self.assertIn("EFT yalnız TL ile yapılır", self.bul("T08"))
        self.assertIn("Rakam (85.200,00) ile yazı (82.500,00) farklı", self.bul("T09"))
        self.assertIn("Alıcı adı/unvanı boş", self.bul("T10"))
        self.assertIn("T01 numaralı talimat da var", self.bul("T11"))
        self.assertIn("havale olarak işlenebilir", self.bul("T12"))
        self.assertIn("Ali Yabancı: müşterinin imza yetkilileri arasında yok", self.bul("T13"))
        self.assertIn("hesap M200 müşterisine ait", self.bul("T14"))
        self.assertIn("hafta sonu", self.bul("T15"))

    def test_kur_ile_limit(self):
        t = main.Talimat("X", date(2026, 10, 8), "SWIFT", "M100", "", "", "A", "", D(20000), "", "USD", "", ["Kemal Örnek"])
        y = main.yetkileri_oku(ORNEK / "imza_yetkileri.csv")["M100"]
        self.assertEqual(main.imza_kontrol(t, y, D(20000) * D("41.20")), [("Hata", "İmza", "Kemal Örnek münferit limiti (500.000,00 TL) "
                                                                                       "tutarın (824.000,00 TL) altında")])
        self.assertEqual(main.imza_kontrol(t, y, D(20000) * D("10")), [])

    def test_excel(self):
        self.assertEqual(self.wb.sheetnames, ["Özet", "Talimatlar", "Bulgular"])
        ws = self.wb["Talimatlar"]
        self.assertEqual(ws["A2"].value, "Hatalı")
        self.assertEqual(ws.cell(1, 15).value, "Kontrol Eden / Onay")


class ListesizTest(unittest.TestCase):
    def test_yalniz_talimat(self):
        with tempfile.TemporaryDirectory() as tmp:
            s = main.calistir(ORNEK / "talimatlar.csv", Path(tmp) / "t.xlsx")
        t = {x.no: x for x in s["talimatlar"]}
        self.assertFalse(any(a == "İmza" for x in t.values() for _, a, _ in x.bulgular))   # yetki listesi yoksa imza kontrolü yok
        self.assertEqual(t["T14"].durum, "Uygun")                                          # hesap listesi yoksa sahiplik bilinmez
        self.assertEqual(t["T07"].durum, "Hatalı")                                         # başka bankaya virman yine yakalanır


if __name__ == "__main__":
    unittest.main()
