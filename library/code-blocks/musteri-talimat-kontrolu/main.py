"""
Müşteri Talimat Kontrolü — Workers / Workless kod bloğu
Bankacılık › Şube Bankacılığı › Gişe Yetkilisi

EFT, FAST, havale ve virman talimatlarını işleme almadan önce kontrol eder:
  - IBAN: biçim, TR IBAN uzunluğu (26), rezerv hane ve mod 97 kontrol hanesi (gönderen ve alıcı).
  - Gönderen hesap müşteriye ait mi, gönderen unvanı hesap unvanıyla uyuşuyor mu (--hesaplar verilirse);
    alıcı hesap bankanın kendi müşterisiyse alıcı unvanı da karşılaştırılır.
  - Tutar: rakamla tutar ile yazıyla tutar aynı mı? Yazı bitişik de olabilir ("Yüzyirmibeşbin TL Elli Kuruş").
  - İşlem türü: havale aynı banka içinde, virman müşterinin kendi hesapları arasında, EFT/FAST başka bankaya
    ve TL olmalı; aynı bankaya EFT, hafta sonu EFT gibi durumlar bilgi olarak işaretlenir.
  - İmza yetkisi (--yetkiler verilirse): imzalayan yetkili mi, yetkisi süresi dolmuş mu, münferit/müşterek
    şartı ve limit sağlanıyor mu?
  - Aynı gün, aynı alıcıya, aynı tutarla verilmiş mükerrer talimatlar.
İnternete bağlanmaz; alıcı unvanının karşı bankadaki IBAN sahibiyle eşleşmesini kontrol edemez.

Kullanım:
    python main.py                                                     # örnek talimatlarla
    python main.py --girdi talimatlar.xlsx --hesaplar hesaplar.xlsx --yetkiler imza_yetkileri.xlsx
    python main.py --girdi talimatlar.xlsx --banka-kodu 00999 --kur USD=41,20 EUR=48,05
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
SIFIR = Decimal(0)
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
TL_ADLARI = {"", "tl", "try", "trl"}
HUKUKI_EKLER = {"as", "a s", "ltd", "sti", "ltd sti", "san", "ve", "tic", "sanayi", "ticaret", "limited", "sirketi", "anonim", "a", "s"}

SUTUNLAR = {
    "no": ("talimat no", "no", "islem no", "sira no"),
    "tarih": ("tarih", "islem tarihi", "talimat tarihi", "valor"),
    "tur": ("tur", "islem turu", "talimat turu"),
    "musteri": ("musteri no", "musteri numarasi", "musteri"),
    "gonderen_unvan": ("gonderen unvan", "gonderen", "gonderen adi", "borclu unvan"),
    "gonderen_iban": ("gonderen iban", "borclu iban", "gonderen hesap", "hesap iban"),
    "alici_unvan": ("alici unvan", "alici", "alici adi", "lehtar", "alici adi soyadi"),
    "alici_iban": ("alici iban", "lehtar iban", "alacakli iban"),
    "tutar": ("tutar", "tutar (rakam)", "tutar rakam", "rakamla tutar"),
    "yazi": ("tutar (yazi)", "tutar yazi", "yaziyla tutar", "yazi ile tutar"),
    "pb": ("para birimi", "doviz", "pb"),
    "aciklama": ("aciklama", "odeme aciklamasi"),
    "imza": ("imzalayanlar", "imzalayan", "imza", "imza sahibi", "imza sahipleri"),
}
HESAP_SUTUN = {"musteri": SUTUNLAR["musteri"], "unvan": ("unvan", "musteri unvani", "ad soyad", "hesap sahibi"), "iban": ("iban",),
               "pb": SUTUNLAR["pb"]}
YETKI_SUTUN = {"musteri": SUTUNLAR["musteri"], "yetkili": ("yetkili", "yetkili adi", "ad soyad", "imza yetkilisi"),
               "sekil": ("yetki sekli", "yetki", "imza sekli"), "limit": ("limit (tl)", "limit", "yetki limiti", "limit tl"),
               "bitis": ("gecerlilik bitis", "bitis tarihi", "gecerlilik", "son gecerlilik")}


def katla(s) -> str:
    return " ".join(str(s or "").translate(_TR).lower().split())


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1 or (s.count(".") == 1 and len(s.split(".")[1]) == 3):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(str(x or "").strip(), f).date()
        except ValueError:
            pass
    return None


def tl(x: Decimal) -> str:
    return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sutunlar: dict, zorunlu: tuple) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:10]):
        b = [katla(x) for x in r]
        k = {a: next((i for i, x in enumerate(b) if x in es), None) for a, es in sutunlar.items()}
        if all(k[z] is not None for z in zorunlu):
            break
    else:
        raise SystemExit(f"{yol.name}: gerekli sütunlar bulunamadı ({', '.join(zorunlu)}). İlk satır: {s[0] if s else '(boş)'}")
    return [{a: (r[i] if i is not None and i < len(r) else None) for a, i in k.items()} for r in s[bi + 1:]]


# ----------------------------------------------------------------------------
# IBAN ve tutar yazısı
# ----------------------------------------------------------------------------

def iban_temizle(x) -> str:
    return re.sub(r"[\s-]", "", str(x or "")).upper()


def iban_kontrol(iban: str) -> str:
    """Boş metin = geçerli. ISO 13616 mod 97; TR IBAN 26 karakter, 10. karakter (rezerv) 0."""
    if not iban:
        return "IBAN boş"
    if not re.fullmatch(r"[A-Z]{2}\d{2}[A-Z0-9]{10,30}", iban):
        return "IBAN biçimi hatalı"
    if iban.startswith("TR"):
        if len(iban) != 26:
            return f"TR IBAN 26 karakter olmalı ({len(iban)})"
        if not iban[2:].isdigit():
            return "TR IBAN'da harf olamaz"
        if iban[9] != "0":
            return "TR IBAN'ın 10. karakteri (rezerv alan) 0 olmalı"
    if int("".join(str(int(c, 36)) for c in iban[4:] + iban[:4])) % 97 != 1:
        return "IBAN kontrol hanesi (mod 97) tutmuyor — yazım hatası olabilir"
    return ""


def banka_kodu(iban: str) -> str | None:
    return iban[4:9] if iban.startswith("TR") and len(iban) == 26 else None


SAYI_SOZ = {"sifir": 0, "bir": 1, "iki": 2, "uc": 3, "dort": 4, "bes": 5, "alti": 6, "yedi": 7, "sekiz": 8, "dokuz": 9, "on": 10, "yirmi": 20,
            "otuz": 30, "kirk": 40, "elli": 50, "altmis": 60, "yetmis": 70, "seksen": 80, "doksan": 90, "yuz": 100, "bin": 1000,
            "milyon": 10 ** 6, "milyar": 10 ** 9}
_SOZ_SIRA = sorted(SAYI_SOZ, key=len, reverse=True)
ANA_BIRIM = {"tl", "try", "lira", "lirasi", "dolar", "usd", "avro", "euro", "eur", "sterlin", "gbp", "frank", "chf"}
ALT_BIRIM = {"kurus", "krs", "kr", "sent", "cent"}


def sozden_sayi(metin: str) -> int | None:
    """Bitişik veya ayrı yazılmış Türkçe sayı ('yüzyirmibeşbin', 'bir milyon iki yüz bin') → tam sayı."""
    s = re.sub(r"\s+", "", metin)
    if not s:
        return None
    toplam, simdiki, i = 0, 0, 0
    while i < len(s):
        soz = next((x for x in _SOZ_SIRA if s.startswith(x, i)), None)
        if soz is None:
            return None
        v = SAYI_SOZ[soz]
        i += len(soz)
        if v == 100:
            simdiki = (simdiki or 1) * 100
        elif v >= 1000:
            toplam += (simdiki or 1) * v
            simdiki = 0
        else:
            simdiki += v
    return toplam + simdiki


def yazidan_tutar(yazi) -> Decimal | None:
    """'Yalnız Yüzyirmibeşbin Türk Lirası Elli Kuruş' → 125000.50. Rakamla yazılmışsa doğrudan okunur."""
    s = katla(yazi)
    if not s:
        return None
    if re.search(r"\d", s):
        return para(re.sub(r"[^\d,.]", "", s))
    kelimeler = [k for k in re.split(r"[^a-z]+", s) if k and k not in {"yalniz", "yalnizca", "turk", "#"}]
    ana, alt, bolum = [], [], 0
    for k in kelimeler:
        if k in ANA_BIRIM:
            bolum = 1
        elif k in ALT_BIRIM:
            bolum = 2
        elif bolum == 0:
            ana.append(k)
        elif bolum == 1:
            alt.append(k)
        else:
            return None                                         # kuruştan sonra sayı olmamalı
    a = sozden_sayi("".join(ana)) if ana else 0
    b = sozden_sayi("".join(alt)) if alt else 0
    if a is None or b is None or (not ana and not alt) or b >= 100:
        return None
    return Decimal(a) + Decimal(b) / 100


def unvan_sade(s) -> set[str]:
    k = re.sub(r"[^a-z0-9 ]+", " ", katla(s))
    return {x for x in k.split() if x not in HUKUKI_EKLER}


def unvan_karsilastir(a, b) -> str:
    """'' = aynı, 'kisa' = biri diğerinin kısaltılmışı, 'farkli'."""
    x, y = unvan_sade(a), unvan_sade(b)
    if x == y:
        return ""
    if x and y and (x <= y or y <= x):
        return "kisa"
    return "farkli"


# ----------------------------------------------------------------------------
# Kontrol
# ----------------------------------------------------------------------------

@dataclass
class Talimat:
    no: str
    tarih: date | None
    tur: str
    musteri: str
    gonderen_unvan: str
    gonderen_iban: str
    alici_unvan: str
    alici_iban: str
    tutar: Decimal | None
    yazi: str
    pb: str
    aciklama: str
    imza: list[str]
    bulgular: list[tuple[str, str, str]] = field(default_factory=list)

    @property
    def durum(self) -> str:
        onemler = {b[0] for b in self.bulgular}
        return "Hatalı" if "Hata" in onemler else "Kontrol" if "Dikkat" in onemler else "Uygun"


def talimatlari_oku(yol: Path) -> list[Talimat]:
    sonuc = []
    for i, r in enumerate(kayitlar(yol, SUTUNLAR, ("tutar",)), 1):
        if all(v in (None, "") for v in r.values()):
            continue
        imza = [x for x in re.split(r"\s*(?:,|;|/|\+|\bve\b)\s*", str(r["imza"] or "").strip()) if x]
        sonuc.append(Talimat(str(r["no"] or f"S{i}").strip(), tarih(r["tarih"]), str(r["tur"] or "").strip(), str(r["musteri"] or "").strip(),
                             str(r["gonderen_unvan"] or "").strip(), iban_temizle(r["gonderen_iban"]), str(r["alici_unvan"] or "").strip(),
                             iban_temizle(r["alici_iban"]), para(r["tutar"]), str(r["yazi"] or "").strip(),
                             str(r["pb"] or "").strip().upper() or "TL", str(r["aciklama"] or "").strip(), imza))
    return sonuc


def hesaplari_oku(yol: Path) -> dict[str, dict]:
    """IBAN → {musteri, unvan, pb}"""
    return {iban_temizle(r["iban"]): {"musteri": str(r["musteri"] or "").strip(), "unvan": str(r["unvan"] or "").strip(),
                                      "pb": str(r["pb"] or "TL").strip().upper()}
            for r in kayitlar(yol, HESAP_SUTUN, ("iban",)) if r["iban"]}


def yetkileri_oku(yol: Path) -> dict[str, list[dict]]:
    y = defaultdict(list)
    for r in kayitlar(yol, YETKI_SUTUN, ("musteri", "yetkili")):
        if r["yetkili"]:
            y[str(r["musteri"]).strip()].append({"ad": str(r["yetkili"]).strip(), "sekil": "Müşterek" if "muster" in katla(r["sekil"]) else "Münferit",
                                                  "limit": para(r["limit"]), "bitis": tarih(r["bitis"])})
    return y


def imza_kontrol(t: Talimat, yetkililer: list[dict], tutar_tl: Decimal | None) -> list[tuple[str, str, str]]:
    if not t.imza:
        return [("Hata", "İmza", "İmzalayan belirtilmemiş")]
    if not yetkililer:
        return [("Hata", "İmza", f"Müşteri {t.musteri or '?'} için imza yetkisi kaydı yok")]
    sonuc, gecerli = [], []
    for ad in t.imza:
        y = next((x for x in yetkililer if katla(x["ad"]) == katla(ad)), None)
        if y is None:
            sonuc.append(("Hata", "İmza", f"{ad}: müşterinin imza yetkilileri arasında yok"))
        elif y["bitis"] and t.tarih and y["bitis"] < t.tarih:
            sonuc.append(("Hata", "İmza", f"{ad}: yetki süresi {y['bitis']:%d.%m.%Y} tarihinde dolmuş"))
        else:
            gecerli.append(y)
    if not gecerli:
        return sonuc
    if tutar_tl is None:
        return sonuc + [("Dikkat", "İmza", "Tutar TL'ye çevrilemedi (kur yok); limit kontrolü yapılmadı")]
    sinirsiz_mi = lambda x: x["limit"] is None  # noqa: E731
    munferit = [x for x in gecerli if x["sekil"] == "Münferit"]
    if any(sinirsiz_mi(x) or x["limit"] >= tutar_tl for x in munferit):
        return sonuc
    if len(gecerli) >= 2:
        if any(sinirsiz_mi(x) for x in gecerli) or max(x["limit"] for x in gecerli) >= tutar_tl:
            return sonuc
        return sonuc + [("Hata", "İmza", f"Müşterek imza limiti ({tl(max(x['limit'] for x in gecerli))} TL) tutarın ({tl(tutar_tl)} TL) altında")]
    y = gecerli[0]
    if y["sekil"] == "Müşterek":
        return sonuc + [("Hata", "İmza", f"{y['ad']} müşterek yetkili; ikinci bir yetkilinin imzası gerekli")]
    return sonuc + [("Hata", "İmza", f"{y['ad']} münferit limiti ({tl(y['limit'])} TL) tutarın ({tl(tutar_tl)} TL) altında")]


def kontrol_et(talimatlar: list[Talimat], hesaplar: dict, yetkiler: dict | None, kurlar: dict[str, Decimal], kendi_banka: str | None) -> None:
    gorulen = {}
    for t in talimatlar:
        b = t.bulgular
        tur = katla(t.tur)
        # IBAN
        for ad, iban in (("Gönderen IBAN", t.gonderen_iban), ("Alıcı IBAN", t.alici_iban)):
            h = iban_kontrol(iban)
            if h:
                b.append(("Hata", ad, h))
        g_iban_ok = not iban_kontrol(t.gonderen_iban)
        a_iban_ok = not iban_kontrol(t.alici_iban)
        # Gönderen hesap ve unvan
        if hesaplar and g_iban_ok:
            h = hesaplar.get(t.gonderen_iban)
            if h is None or (t.musteri and h["musteri"] != t.musteri):
                b.append(("Hata", "Gönderen hesap", f"Gönderen IBAN müşteri {t.musteri or '?'} hesapları arasında yok"
                          + (f" (hesap {h['musteri']} müşterisine ait)" if h else "")))
            else:
                k = unvan_karsilastir(t.gonderen_unvan, h["unvan"])
                if k == "farkli":
                    b.append(("Dikkat", "Gönderen unvan", f"Talimattaki unvan ({t.gonderen_unvan}) hesap unvanıyla ({h['unvan']}) uyuşmuyor"))
                elif k == "kisa":
                    b.append(("Bilgi", "Gönderen unvan", f"Unvan kısaltılmış yazılmış; hesap unvanı: {h['unvan']}"))
                pb = lambda x: "tl" if katla(x) in TL_ADLARI else katla(x)  # noqa: E731
                if pb(h["pb"]) != pb(t.pb):
                    b.append(("Dikkat", "Para birimi", f"Gönderen hesap {h['pb']}, talimat {t.pb}"))
        # Alıcı
        if not t.alici_unvan:
            b.append(("Hata", "Alıcı unvan", "Alıcı adı/unvanı boş"))
        alici_hesap = hesaplar.get(t.alici_iban) if hesaplar and a_iban_ok else None
        if alici_hesap and t.alici_unvan and unvan_karsilastir(t.alici_unvan, alici_hesap["unvan"]) == "farkli":
            b.append(("Hata", "Alıcı unvan", f"Alıcı IBAN bankamızda {alici_hesap['unvan']} adına kayıtlı; talimatta: {t.alici_unvan}"))
        # Tutar
        if t.tutar is None or t.tutar <= 0:
            b.append(("Hata", "Tutar", "Rakamla tutar okunamadı veya sıfır"))
        if not t.yazi:
            b.append(("Dikkat", "Tutar", "Yazıyla tutar yok"))
        elif t.tutar is not None:
            y = yazidan_tutar(t.yazi)
            if y is None:
                b.append(("Dikkat", "Tutar", f"Yazıyla tutar okunamadı: '{t.yazi}' — elle kontrol edin"))
            elif y != t.tutar:
                b.append(("Hata", "Tutar", f"Rakam ({tl(t.tutar)}) ile yazı ({tl(y)}) farklı"))
        # Tür
        ayni_banka = a_iban_ok and kendi_banka and banka_kodu(t.alici_iban) == kendi_banka
        kendi_hesabi = bool(alici_hesap and t.musteri and alici_hesap["musteri"] == t.musteri)
        if tur == "virman":
            if a_iban_ok and kendi_banka and not ayni_banka:
                b.append(("Hata", "İşlem türü", "Virman yalnız bankamızdaki hesaplar arasında yapılır; alıcı başka bankada (EFT/FAST olmalı)"))
            elif hesaplar and a_iban_ok and not kendi_hesabi:
                b.append(("Hata", "İşlem türü", "Virman müşterinin kendi hesapları arasında yapılır; alıcı başka müşteri (havale olmalı)"))
        elif tur == "havale":
            if a_iban_ok and kendi_banka and not ayni_banka:
                b.append(("Hata", "İşlem türü", "Havale banka içidir; alıcı başka bankada (EFT/FAST olmalı)"))
            elif kendi_hesabi:
                b.append(("Bilgi", "İşlem türü", "Alıcı müşterinin kendi hesabı: virman olarak işlenebilir"))
        elif tur in ("eft", "fast"):
            if katla(t.pb) not in TL_ADLARI:
                b.append(("Hata", "İşlem türü", f"{t.tur.upper()} yalnız TL ile yapılır; {t.pb} transfer için döviz transferi (SWIFT) gerekir"))
            if ayni_banka:
                b.append(("Bilgi", "İşlem türü", "Alıcı bankamızda: " + ("virman" if kendi_hesabi else "havale") + " olarak işlenebilir"))
            if tur == "eft" and t.tarih and t.tarih.weekday() >= 5:
                b.append(("Bilgi", "Tarih", f"{t.tarih:%d.%m.%Y} hafta sonu; EFT sistemi iş günlerinde çalışır"))
        elif not tur:
            b.append(("Dikkat", "İşlem türü", "Tür belirtilmemiş"))
        # İmza
        if yetkiler is not None:
            kur = Decimal(1) if katla(t.pb) in TL_ADLARI else kurlar.get(t.pb)
            b += imza_kontrol(t, yetkiler.get(t.musteri, []), None if t.tutar is None or kur is None else t.tutar * kur)
        # Mükerrer
        anahtar = (t.musteri or t.gonderen_iban, t.alici_iban, t.tutar, t.tarih)
        if anahtar in gorulen:
            b.append(("Dikkat", "Mükerrer", f"Aynı gün, aynı alıcı ve tutarla {gorulen[anahtar]} numaralı talimat da var"))
        else:
            gorulen[anahtar] = t.no


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Hatalı": "FDE2E1", "Kontrol": "FFF4CE", "Uygun": "E3F5E1", "Hata": "FDE2E1", "Dikkat": "FFF4CE", "Bilgi": "E8F0FE"}
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def rapor_yaz(cikti: Path, talimatlar: list[Talimat], kendi_banka: str | None, kaynaklar: list[str]) -> None:
    wb = Workbook()
    o = wb.active
    o.title = "Özet"
    o.column_dimensions["A"].width, o.column_dimensions["B"].width = 34, 60
    durum = Counter(t.durum for t in talimatlar)
    for etiket, deger in [("Talimat", len(talimatlar)), ("Hatalı", durum["Hatalı"]), ("Kontrol gerekli", durum["Kontrol"]),
                          ("Uygun", durum["Uygun"]), ("Bankamız kodu", kendi_banka or "belirlenemedi"), ("Kullanılan listeler", ", ".join(kaynaklar))]:
        o.append([etiket, deger])
        o.cell(o.max_row, 1).font = Font(bold=True)
        if etiket in RENK:
            o.cell(o.max_row, 1).fill = PatternFill("solid", fgColor=RENK[etiket])
    o.append([])
    _baslik(o, ["Kontrol alanı", "Hata", "Dikkat", "Bilgi"], (34, 10, 10, 10))
    alan = defaultdict(Counter)
    for t in talimatlar:
        for onem, a, _ in t.bulgular:
            alan[a][onem] += 1
    for a, c in sorted(alan.items(), key=lambda i: (-i[1]["Hata"], -i[1]["Dikkat"], i[0])):
        o.append([a, c["Hata"] or None, c["Dikkat"] or None, c["Bilgi"] or None])
    o.append([])
    o.append(["Not", "Ön kontroldür. İmza karşılaştırması (imza örneği), kimlik/yetki belgelerinin aslı ve alıcı unvanının karşı bankadaki "
                     "IBAN sahibiyle eşleşmesi bu kodla kontrol edilemez; bankanızın prosedürüne göre yapılmalıdır."])
    o.cell(o.max_row, 2).alignment = UST

    ws = wb.create_sheet("Talimatlar")
    _baslik(ws, ["Durum", "Talimat No", "Tarih", "Tür", "Müşteri No", "Gönderen Unvan", "Gönderen IBAN", "Alıcı Unvan", "Alıcı IBAN", "Tutar",
                 "PB", "Tutar (Yazı)", "İmzalayanlar", "Bulgular", "Kontrol Eden / Onay"],
            (10, 10, 11, 8, 10, 28, 30, 28, 30, 15, 6, 34, 26, 60, 18))
    for t in sorted(talimatlar, key=lambda t: ({"Hatalı": 0, "Kontrol": 1, "Uygun": 2}[t.durum], t.no)):
        ws.append([t.durum, t.no, t.tarih, t.tur, t.musteri, t.gonderen_unvan, t.gonderen_iban, t.alici_unvan, t.alici_iban,
                   None if t.tutar is None else float(t.tutar), t.pb, t.yazi, ", ".join(t.imza),
                   "\n".join(f"[{o_}] {a}: {m}" for o_, a, m in t.bulgular), ""])
        ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=RENK[t.durum])
        ws.cell(ws.max_row, 3).number_format = "DD.MM.YYYY"
        ws.cell(ws.max_row, 10).number_format = "#,##0.00"
        ws.cell(ws.max_row, 15).fill = PatternFill("solid", fgColor="FFF4CE")
        for h in ws[ws.max_row]:
            h.alignment = UST
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    bs = wb.create_sheet("Bulgular")
    _baslik(bs, ["Önem", "Talimat No", "Alan", "Açıklama"], (9, 11, 18, 100))
    for t in talimatlar:
        for onem, a, m in t.bulgular:
            bs.append([onem, t.no, a, m])
            bs.cell(bs.max_row, 1).fill = PatternFill("solid", fgColor=RENK[onem])
    bs.auto_filter.ref = bs.dimensions
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


# ----------------------------------------------------------------------------
# Akış
# ----------------------------------------------------------------------------

def calistir(girdi: Path, cikti: Path, hesap_yolu: Path | None = None, yetki_yolu: Path | None = None, kendi_banka: str | None = None,
             kurlar: dict[str, Decimal] | None = None) -> dict:
    talimatlar = talimatlari_oku(girdi)
    hesaplar = hesaplari_oku(hesap_yolu) if hesap_yolu else {}
    yetkiler = yetkileri_oku(yetki_yolu) if yetki_yolu else None
    if not kendi_banka:                                         # en sık görülen gönderen banka kodu
        kodlar = Counter(banka_kodu(i) for i in [*hesaplar, *(t.gonderen_iban for t in talimatlar)] if not iban_kontrol(i) and banka_kodu(i))
        kendi_banka = kodlar.most_common(1)[0][0] if kodlar else None
    kontrol_et(talimatlar, hesaplar, yetkiler, kurlar or {}, kendi_banka)
    kaynaklar = ["talimatlar"] + (["hesaplar"] if hesaplar else []) + (["imza yetkileri"] if yetkiler is not None else [])
    rapor_yaz(cikti, talimatlar, kendi_banka, kaynaklar)
    return {"talimatlar": talimatlar, "banka": kendi_banka}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="EFT, FAST, havale ve virman talimatlarını IBAN, unvan, rakam-yazı tutar ve imza yetkisi açısından kontrol eder.")
    p.add_argument("--girdi", type=Path, default=ORNEK / "talimatlar.csv", help="Talimatlar (.xlsx/.csv)")
    p.add_argument("--hesaplar", type=Path, help="Müşteri hesapları: Müşteri No, Unvan, IBAN, Para Birimi")
    p.add_argument("--yetkiler", type=Path, help="İmza yetkileri: Müşteri No, Yetkili, Yetki Şekli (Münferit/Müşterek), Limit (TL), Geçerlilik Bitiş")
    p.add_argument("--banka-kodu", help="Bankanızın 5 haneli banka kodu (verilmezse gönderen IBAN'lardan bulunur)")
    p.add_argument("--kur", nargs="*", default=[], metavar="PB=KUR", help="Döviz talimatlarında limit kontrolü için kur, ör. USD=41,20")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "talimat_kontrolu.xlsx")
    a = p.parse_args(argv)
    hesap, yetki = a.hesaplar, a.yetkiler
    if a.girdi == ORNEK / "talimatlar.csv":
        hesap, yetki = hesap or ORNEK / "hesaplar.csv", yetki or ORNEK / "imza_yetkileri.csv"
    kurlar = {}
    for x in a.kur:
        ad, _, v = x.partition("=")
        if para(v) is None:
            print(f"[X] --kur 'USD=41,20' biçiminde olmalı: {x}")
            return 2
        kurlar[ad.strip().upper()] = para(v)
    s = calistir(a.girdi, a.cikti, hesap, yetki, a.banka_kodu, kurlar)
    d = Counter(t.durum for t in s["talimatlar"])
    print(f"[OK] {len(s['talimatlar'])} talimat · hatalı {d['Hatalı']} · kontrol {d['Kontrol']} · uygun {d['Uygun']} · banka kodu {s['banka'] or '?'}")
    for t in s["talimatlar"]:
        for o, al, m in t.bulgular:
            if o == "Hata":
                print(f"[X] {t.no} {al}: {m}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
