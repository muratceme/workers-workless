"""
Reçete Maliyet Hesaplama — Workers / Workless kod bloğu
Ar-Ge ve Ürün Geliştirme › Ar-Ge Uzmanı

Reçete bileşen miktarları, hammadde fiyatları ve fire oranlarından kilogram ve ambalaj başına hammadde maliyetini hesaplar.
  Bileşen oranı %     = bileşen miktarı ÷ reçete toplamı × 100
  Brüt ihtiyaç        = miktar ÷ (1 − fire % ÷ 100)                         (ayıklama, eleme, hazırlama firesi: reçeteye net miktar girer)
  Bileşen maliyeti    = brüt ihtiyaç × birim fiyat
  Mamul miktarı       = reçete toplamı × proses verimi % ÷ 100              (pişirme, buharlaşma, hat kaybı; verilmezse %100)
  Kg maliyeti         = Σ bileşen maliyeti ÷ mamul miktarı
  Ambalaj başına      = kg maliyeti × net ağırlık (g) ÷ 1000 + ambalaj malzemesi maliyeti
  Maliyet payı %      = bileşen maliyeti ÷ Σ bileşen maliyeti × 100
Yalnız hammadde ve ambalaj malzemesi hesaplanır; işçilik, enerji, amortisman, genel gider ve kâr dahil değildir. Miktarlar ve fiyatlar aynı birimde (ör. kg ve TL/kg) olmalıdır.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


RECETE = {"urun": ("urun", "recete", "urun adi", "mamul"), "bilesen": ("bilesen", "hammadde", "malzeme"), "miktar": ("miktar", "miktar kg", "oran", "kg"), "fire": ("fire", "fire yuzde", "fire orani")}
FIYAT = {"bilesen": RECETE["bilesen"], "fiyat": ("birim fiyat", "birim fiyat tl kg", "fiyat", "alis fiyati", "tl kg")}
URUN = {"urun": RECETE["urun"], "verim": ("proses verimi", "verim", "proses verimi yuzde"), "net": ("ambalaj net", "ambalaj net g", "net agirlik", "net agirlik g", "net gramaj", "gramaj"),
        "ambalaj": ("ambalaj maliyeti", "ambalaj malzemesi maliyeti", "ambalaj")}
Z, YUZ = Decimal(0), Decimal(100)


def olcu(x) -> Decimal | None:
    """Miktar ve oran: 12,5 = 12.5; 1.250,5 = bin iki yüz elli buçuk. Tek başına nokta ondalık ayırıcıdır."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    else:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def oku(recete_yol: Path, fiyat_yol: Path, urun_yol: Path | None = None) -> tuple[list[dict], dict, dict, list[dict]]:
    uy, recete, fiyat, urun = [], [], {}, {}
    for r in kayitlar(recete_yol, RECETE, ("urun", "bilesen", "miktar")):
        u, b = metin(r["urun"]), metin(r["bilesen"])
        if not u and not b:
            continue
        m, f = olcu(r["miktar"]), olcu(r["fire"]) if metin(r["fire"]) else Z
        if not u or not b or m is None or m <= 0 or f is None or not 0 <= f < 100:
            uy.append(uyari("Yüksek", "Reçete satırı okunamadı", f"{u or '?'} · {b or '?'} (satır {r['_satir']})", f"Miktar '{metin(r['miktar'])}', fire '{metin(r['fire'])}'. Miktar pozitif, fire 0–100 arasında olmalı. Satır alınmadı; ürün maliyeti eksik."))
            continue
        recete.append({"urun": u, "bilesen": b, "miktar": m, "fire": f})
    for r in kayitlar(fiyat_yol, FIYAT, ("bilesen", "fiyat")):
        b = metin(r["bilesen"])
        if not b:
            continue
        f = para(r["fiyat"])
        if f is None or f < 0:
            uy.append(uyari("Yüksek", "Fiyat okunamadı", b, f"Birim fiyat '{metin(r['fiyat'])}'. Bu bileşen fiyatsız sayıldı."))
            continue
        if katla(b) in fiyat:
            uy.append(uyari("Orta", "Mükerrer fiyat", b, f"Bileşen fiyat listesinde birden fazla kez var. İlk fiyat ({fiyat[katla(b)]:f}) kullanıldı."))
            continue
        fiyat[katla(b)] = f
    if urun_yol:
        for r in kayitlar(urun_yol, URUN, ("urun",)):
            u = metin(r["urun"])
            if not u:
                continue
            v, n, a = (olcu(r["verim"]) if metin(r["verim"]) else YUZ), (olcu(r["net"]) if metin(r["net"]) else None), (para(r["ambalaj"]) if metin(r["ambalaj"]) else Z)
            if v is None or v <= 0 or (metin(r["net"]) and (n is None or n <= 0)) or a is None or a < 0:
                uy.append(uyari("Orta", "Ürün bilgisi okunamadı", u, f"Proses verimi '{metin(r['verim'])}', net ağırlık '{metin(r['net'])}', ambalaj maliyeti '{metin(r['ambalaj'])}'. Ürün %100 verimle, ambalajsız hesaplandı."))
                continue
            urun[katla(u)] = {"verim": v, "net": n, "ambalaj": a}
    return recete, fiyat, urun, uy


def hesapla(recete: list[dict], fiyat: dict, urun: dict | None = None) -> dict:
    if not recete:
        raise ValueError("Hesaplanacak reçete satırı yok.")
    urun, uy, grup = urun or {}, [], {}
    for s in recete:
        grup.setdefault(katla(s["urun"]), {"urun": s["urun"], "satirlar": []})["satirlar"].append(s)
    urunler, satirlar = [], []
    for uk, g in grup.items():
        toplam = sum(s["miktar"] for s in g["satirlar"])
        bilgi = urun.get(uk, {"verim": YUZ, "net": None, "ambalaj": Z})
        ara, maliyet, eksik = [], Z, []
        for s in g["satirlar"]:
            f = fiyat.get(katla(s["bilesen"]))
            brut = s["miktar"] / (1 - s["fire"] / 100)
            m = brut * f if f is not None else None
            if f is None:
                eksik.append(s["bilesen"])
            else:
                maliyet += m
            ara.append({"urun": g["urun"], "bilesen": s["bilesen"], "miktar": s["miktar"], "oran": yuvarla(s["miktar"] / toplam * 100), "fire": s["fire"], "brut": yuvarla(brut, "0.001"), "fiyat": f, "ham": m})
        for a in ara:
            m = a.pop("ham")
            satirlar.append(a | {"maliyet": yuvarla(m) if m is not None else None, "pay": yuvarla(m / maliyet * 100) if m is not None and maliyet else None})
        if eksik:
            uy.append(uyari("Yüksek", "Fiyatı olmayan bileşen", g["urun"], f"Fiyat listesinde yok: {', '.join(eksik)}. Ürün maliyeti bu bileşenler olmadan hesaplandı; gerçek maliyet daha yüksektir."))
        mamul = toplam * bilgi["verim"] / 100
        kg = maliyet / mamul
        adet_ham = kg * bilgi["net"] / 1000 if bilgi["net"] else None
        urunler.append({"urun": g["urun"], "bilesen": len(ara), "toplam": toplam, "maliyet": yuvarla(maliyet), "verim": bilgi["verim"], "mamul": yuvarla(mamul, "0.001"), "kg": yuvarla(kg), "net": bilgi["net"],
                        "ambalaj": bilgi["ambalaj"] if bilgi["net"] else None, "adet_ham": yuvarla(adet_ham) if adet_ham is not None else None,
                        "adet": yuvarla(adet_ham + bilgi["ambalaj"]) if adet_ham is not None else None, "eksik": eksik})
    for uk in urun:
        if uk not in grup:
            uy.append(uyari("Bilgi", "Reçetesi olmayan ürün", uk, "Ürün bilgisi dosyasında var ama reçete dosyasında satırı yok."))
    return {"urunler": urunler, "satirlar": satirlar, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ürün Maliyeti"
    _baslik(ws, ["Ürün", "Bileşen Sayısı", "Reçete Toplamı", "Hammadde Maliyeti", "Proses Verimi %", "Mamul Miktarı", "Kg Maliyeti", "Ambalaj Net (g)", "Ambalaj Başına Hammadde", "Ambalaj Malzemesi", "Ambalaj Başına Toplam", "Durum"],
            (30, 9, 12, 15, 10, 12, 12, 10, 14, 12, 14, 18))
    for u in r["urunler"]:
        _satir(ws, [u["urun"], u["bilesen"], u["toplam"], u["maliyet"], u["verim"], u["mamul"], u["kg"], u["net"], u["adet_ham"], u["ambalaj"], u["adet"], "Eksik fiyat" if u["eksik"] else "Tam"],
               para_s=(3, 4, 5, 6, 7, 9, 10, 11), dolgu={12: KIRMIZI if u["eksik"] else YESIL})
    ws.append([])
    ws.append(["Yalnız hammadde ve ambalaj malzemesi maliyetidir; işçilik, enerji, amortisman, genel gider ve kâr dahil değildir. Kg maliyeti = hammadde maliyeti ÷ (reçete toplamı × proses verimi)."])
    ws = wb.create_sheet("Bileşen Detayı")
    _baslik(ws, ["Ürün", "Bileşen", "Miktar", "Oran %", "Fire %", "Brüt İhtiyaç", "Birim Fiyat", "Maliyet", "Maliyet Payı %"], (30, 28, 11, 9, 8, 12, 12, 14, 11))
    for s in r["satirlar"]:
        _satir(ws, [s["urun"], s["bilesen"], s["miktar"], s["oran"], s["fire"] or None, s["brut"], s["fiyat"], s["maliyet"], s["pay"]], para_s=(3, 4, 5, 7, 8, 9), dolgu={7: KIRMIZI if s["fiyat"] is None else None})
        ws.cell(ws.max_row, 6).number_format = "#,##0.000"
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(recete_yol: Path, fiyat_yol: Path, cikti: Path, urun_yol: Path | None = None) -> dict:
    recete, fiyat, urun, uy = oku(recete_yol, fiyat_yol, urun_yol)
    r = hesapla(recete, fiyat, urun)
    uy += r["uyarilar"]
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Reçete, hammadde fiyatı ve fire oranlarından kg ve ambalaj başına maliyeti hesaplar.")
    p.add_argument("--recete", type=Path, default=ORNEK / "recete.csv", help="Ürün, Bileşen, Miktar, Fire %%")
    p.add_argument("--fiyatlar", type=Path, help="Bileşen, Birim Fiyat")
    p.add_argument("--urunler", type=Path, help="İsteğe bağlı: Ürün, Proses Verimi %%, Ambalaj Net (g), Ambalaj Maliyeti")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "recete_maliyet.xlsx")
    a = p.parse_args(argv)
    ornek = a.recete == ORNEK / "recete.csv"
    fiyat, urun = a.fiyatlar or (ORNEK / "fiyatlar.csv" if ornek else None), a.urunler or (ORNEK / "urunler.csv" if ornek else None)
    if fiyat is None:
        print("[X] --fiyatlar dosyası verilmeli")
        return 2
    for y in (a.recete, fiyat, urun):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.recete, fiyat, a.cikti, urun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    for u in r["urunler"]:
        ek = f" · ambalaj başına {u['adet']:,.2f}" if u["adet"] is not None else ""
        print(f"[{'!' if u['eksik'] else 'OK'}] {u['urun']}: kg maliyeti {u['kg']:,.2f}{ek}" + (f" · fiyatı olmayan bileşen: {', '.join(u['eksik'])}" if u["eksik"] else ""))
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
