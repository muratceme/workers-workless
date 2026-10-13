import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def s(urun, bilesen, miktar, fire=0):
    return {"urun": urun, "bilesen": bilesen, "miktar": D(miktar), "fire": D(fire)}


class HesapTesti(unittest.TestCase):
    def test_temel(self):
        r = main.hesapla([s("U", "A", 60, 20), s("U", "B", 40)], {"a": D(10), "b": D(5)}, {"u": {"verim": D(80), "net": D(250), "ambalaj": D("1.50")}})
        self.assertEqual([(x["bilesen"], x["oran"], x["brut"], x["fiyat"], x["maliyet"], x["pay"]) for x in r["satirlar"]], [("A", D("60.00"), D("75.000"), D(10), D("750.00"), D("78.95")), ("B", D("40.00"), D("40.000"), D(5), D("200.00"), D("21.05"))])
        u = r["urunler"][0]
        self.assertEqual((u["toplam"], u["maliyet"], u["mamul"], u["kg"], u["adet_ham"], u["adet"], u["eksik"], r["uyarilar"]), (D(100), D("950.00"), D("80.000"), D("11.88"), D("2.97"), D("4.47"), [], []))

    def test_varsayilan_ve_eksik_fiyat(self):
        r = main.hesapla([s("U", "A", 3), s("u", "B", 1), s("V", "A", 2, 50), s("V", "C", 2)], {"a": D(9), "c": D(0)}, {"x": {"verim": D(90), "net": None, "ambalaj": D(0)}, "v": {"verim": D(50), "net": None, "ambalaj": D(0)}})
        self.assertEqual([(u["urun"], u["bilesen"], u["toplam"], u["maliyet"], u["verim"], u["kg"], u["net"], u["adet"], u["eksik"]) for u in r["urunler"]],
                         [("U", 2, D(4), D("27.00"), D(100), D("6.75"), None, None, ["B"]), ("V", 2, D(4), D("36.00"), D(50), D("18.00"), None, None, [])])
        self.assertEqual([(x["bilesen"], x["maliyet"], x["pay"]) for x in r["satirlar"]], [("A", D("27.00"), D("100.00")), ("B", None, None), ("A", D("36.00"), D("100.00")), ("C", D("0.00"), D("0.00"))])
        self.assertEqual(sorted((x["tur"], x["kim"]) for x in r["uyarilar"]), [("Fiyatı olmayan bileşen", "U"), ("Reçetesi olmayan ürün", "x")])
        with self.assertRaises(ValueError):
            main.hesapla([], {})


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "r.xlsx"
            r = main.calistir(ORNEK / "recete.csv", ORNEK / "fiyatlar.csv", cikti, ORNEK / "urunler.csv")
            self.assertEqual([(u["urun"], u["toplam"], u["maliyet"], u["mamul"], u["kg"], u["adet_ham"], u["adet"], u["eksik"]) for u in r["urunler"]],
                             [("Örnek Domates Sosu", D(100), D("3176.27"), D("95.000"), D("33.43"), D("16.72"), D("23.22"), []),
                              ("Örnek Acı Sos", D(100), D("1863.48"), D("96.000"), D("19.41"), D("4.85"), D("9.85"), ["Acı Biber Püresi"])])
            self.assertEqual([(x["bilesen"], x["brut"], x["maliyet"]) for x in r["satirlar"]][:3], [("Domates Salçası", D("40.816"), D("2448.98")), ("Su", D("50.000"), D("2.50")), ("Ayçiçek Yağı", D("5.051"), D("378.79"))])
            self.assertEqual([(u["tur"], u["kim"]) for u in r["uyarilar"]], [("Fiyatı olmayan bileşen", "Örnek Acı Sos")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Ürün Maliyeti", "Bileşen Detayı", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            ry, fy, uy_ = Path(tmp) / "r.csv", Path(tmp) / "f.csv", Path(tmp) / "u.csv"
            ry.write_text("Reçete;Hammadde;Miktar;Fire\nU;A;12,5;%2,5\nU;B;0;\nU;C;5;100\n;;;\nU;;5;\n", encoding="utf-8")
            fy.write_text("Hammadde;Fiyat\nA;1.250,50\nA;5\nB;-1\n", encoding="utf-8")
            uy_.write_text("Ürün;Verim;Gramaj;Ambalaj\nU;95;500;2,75\nV;0;;\nW;;;\n", encoding="utf-8")
            recete, fiyat, urun, uy = main.oku(ry, fy, uy_)
            self.assertEqual((recete, fiyat, urun, sorted(u["tur"] for u in uy)),
                             ([{"urun": "U", "bilesen": "A", "miktar": D("12.5"), "fire": D("2.5")}], {"a": D("1250.50")}, {"u": {"verim": D(95), "net": D(500), "ambalaj": D("2.75")}, "w": {"verim": D(100), "net": None, "ambalaj": D(0)}},
                              ["Fiyat okunamadı", "Mükerrer fiyat"] + ["Reçete satırı okunamadı"] * 3 + ["Ürün bilgisi okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--recete", str(ry), "--fiyatlar", str(fy), "--cikti", c]
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(arg + ["--urunler", str(uy_)]), 0)
            self.assertEqual(main.main(["--recete", str(ry), "--cikti", c]), 2)
            self.assertEqual(main.main(arg + ["--urunler", str(Path(tmp) / "yok.csv")]), 1)
            ry.write_text("Ürün;Bileşen\nU;A\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
