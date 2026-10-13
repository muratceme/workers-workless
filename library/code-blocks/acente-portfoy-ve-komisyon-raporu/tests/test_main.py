import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def p(no, sirket, brans, tanzim, brut, net=None, komisyon=None, oran=None, tahsil=None):
    c = lambda x: None if x is None else D(str(x))   # noqa: E731
    return {"no": no, "sirket": sirket, "brans": brans, "tanzim": tanzim, "net": c(net), "brut": c(brut), "komisyon": c(komisyon), "oran": c(oran), "tahsil": c(tahsil)}


def veri():
    return [p("1", "Ş1", "Kasko", date(2026, 1, 15), 1200, 1000, 150, tahsil=1200), p("2", "ş1", "Trafik", date(2026, 1, 20), 600, 500, oran=10, tahsil=300), p("3", "Ş2", "Kasko", date(2026, 2, 5), 2000, komisyon=200, tahsil=0),
            p("4", "Ş1", "KASKO", date(2026, 2, 10), -240, -200, -30, tahsil=-240), p("5", "Ş2", "Yangın", None, 480, 400, tahsil=480)]


def k(s):
    return (s["police"], s["iptal"], s["brut"], s["pay"], s["net"], s["komisyon"], s["kom_oran"], s["tahsil"], s["alacak"], s["tahsilat"])


class RaporTesti(unittest.TestCase):
    def test_gostergeler(self):
        r = main.raporla(veri())
        self.assertEqual({x["no"]: (x["net"], x["komisyon"], x["kom_oran"], x["tahsil"], x["alacak"], x["komisyonsuz"]) for x in r["policeler"]},
                         {"1": (D(1000), D(150), D("15.00"), D(1200), D(0), False), "2": (D(500), D("50.00"), D("10.00"), D(300), D(300), False), "3": (D(2000), D(200), D("10.00"), D(0), D(2000), False),
                          "4": (D(-200), D(-30), None, D(-240), D(0), False), "5": (D(400), D(0), D("0.00"), D(480), D(0), True)})
        self.assertEqual([(s["sirket"],) + k(s) for s in r["sirketler"]],
                         [("Ş2", 2, 0, D(2480), D("61.39"), D(2400), D(200), D("8.33"), D(480), D(2000), D("19.35")), ("Ş1", 2, 1, D(1560), D("38.61"), D(1300), D("170.00"), D("13.08"), D(1260), D(300), D("80.77"))])
        self.assertEqual(k(r["toplam"]), (4, 1, D(4040), D("100.00"), D(3700), D("370.00"), D("10.00"), D(1740), D(2300), D("43.07")))
        self.assertEqual([(s["brans"], s["brut"], s["net"], s["komisyon"], s["kom_oran"], s["tahsil"]) for s in r["branslar"]],
                         [("Kasko", D(2960), D(2800), D(320), D("11.43"), D(960)), ("Trafik", D(600), D(500), D("50.00"), D("10.00"), D(300)), ("Yangın", D(480), D(400), D(0), D("0.00"), D(480))])
        self.assertEqual([(s["sirket"], s["brans"], s["brut"], s["police"], s["iptal"]) for s in r["detay"]], [("Ş2", "Kasko", D(2000), 1, 0), ("Ş1", "Kasko", D(960), 1, 1), ("ş1", "Trafik", D(600), 1, 0), ("Ş2", "Yangın", D(480), 1, 0)])
        self.assertEqual(r["aylar"], [{"ay": "2026-01", "brut": D(1800), "komisyon": D("200.00"), "police": 2}, {"ay": "2026-02", "brut": D(1760), "komisyon": D(170), "police": 1}])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]), [("Bilgi", "Portföy tek şirkette yoğunlaşıyor", "Ş2"), ("Orta", "Komisyonu yazılmamış poliçe", "Ş2"), ("Orta", "Tahsilat oranı düşük", "Ş2")])

    def test_esikler(self):
        u = lambda **kw: sorted((x["tur"], x["kim"]) for x in main.raporla(veri(), **kw)["uyarilar"] if x["tur"] != "Komisyonu yazılmamış poliçe")   # noqa: E731
        self.assertEqual(u(yogunluk=D(38), tahsilat=D(81)), [("Portföy tek şirkette yoğunlaşıyor", "Ş1"), ("Portföy tek şirkette yoğunlaşıyor", "Ş2"), ("Tahsilat oranı düşük", "Ş1"), ("Tahsilat oranı düşük", "Ş2")])
        self.assertEqual(u(yogunluk=D("61.39"), tahsilat=D("19.35")), [])
        r = main.raporla([p("1", "Ş", "K", None, -100, tahsil=-100)])
        self.assertEqual((k(r["toplam"]), r["aylar"], r["uyarilar"][0]["tur"]), ((0, 1, D(-100), None, D(-100), D(0), None, D(-100), D(0), None), [], "Komisyonu yazılmamış poliçe"))
        with self.assertRaises(ValueError):
            main.raporla([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "policeler.csv", cikti)
            self.assertEqual([(s["sirket"],) + k(s) for s in r["sirketler"]],
                             [("Deneme Sigorta", 3, 0, D(157500), D("50.71"), D(150000), D(23500), D("15.67"), D(105000), D(52500), D("66.67")),
                              ("Örnek Sigorta", 5, 1, D(142600), D("45.91"), D(137000), D("19950.00"), D("14.56"), D(100600), D(42000), D("70.55")),
                              ("Numune Sigorta", 1, 0, D(10500), D("3.38"), D(10000), D("1000.00"), D("10.00"), D(10500), D(0), D("100.00"))])
            self.assertEqual(k(r["toplam"]), (9, 1, D(310600), D("100.00"), D(297000), D("44450.00"), D("14.97"), D(216100), D(94500), D("69.58")))
            self.assertEqual([(s["brans"], s["brut"]) for s in r["branslar"]], [("Kasko", D(157500)), ("Yangın", D(105000)), ("Sağlık", D(25000)), ("Trafik", D(23100))])
            self.assertEqual([(a["ay"], a["police"], a["brut"]) for a in r["aylar"]], [("2026-07", 3, D(86100)), ("2026-08", 3, D(161500)), ("2026-09", 3, D(63000))])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Orta", "Komisyonu yazılmamış poliçe", "Deneme Sigorta"), ("Orta", "Tahsilat oranı düşük", "Deneme Sigorta"), ("Bilgi", "Portföy tek şirkette yoğunlaşıyor", "Deneme Sigorta")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Şirketler", "Branşlar", "Şirket × Branş", "Aylar", "Poliçeler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "p.csv"
            y.write_text("Poliçe;Şirket;Ürün;Tarih;Net Prim;Prim;Komisyon Tutarı;Komisyon Yüzdesi;Tahsilat\nP1;A;Kasko;x;1.000;1.050,50;;12,5;500\n;B;;;;200;;;\nP3;A;;;;x;;;\nP4;A;;;;10;;150;\nP5;A;;;y;10;;;\nP6;;;;;10;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["sirket"], x["brans"], x["tanzim"], x["net"], x["brut"], x["komisyon"], x["oran"], x["tahsil"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("P1", "A", "Kasko", None, D(1000), D("1050.50"), None, D("12.5"), D(500)), ("Satır 3", "B", "(branş yok)", None, None, D(200), None, None, None)],
                              [("Tarih okunamadı", "P1"), ("Poliçe okunamadı", "P3"), ("Poliçe okunamadı", "P4"), ("Poliçe okunamadı", "P5")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--policeler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--yogunluk", "40", "--tahsilat", "85"]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--yogunluk", "0"], ["--tahsilat", "101"], ["--tahsilat", "x"])], [2] * 3)
            self.assertEqual(main.main(["--policeler", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Sigorta Şirketi;Branş\nA;B\n", "Sigorta Şirketi;Brüt Prim\n;5\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
