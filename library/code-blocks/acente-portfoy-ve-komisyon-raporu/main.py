"""
Acente Portföy ve Komisyon Raporu — Workers / Workless kod bloğu
Teknik ve Operasyon › Acente Yöneticisi

Acentenin üretimini sigorta şirketi ve branş bazında toplar; komisyon gelirini ve tahsilat durumunu raporlar.
  Komisyon             = dosyadaki komisyon tutarı; boşsa net prim × komisyon oranı ÷ 100
  Komisyon oranı %     = Σ komisyon ÷ Σ net prim × 100            (net prim yazılmamışsa brüt prim esas alınır)
  Açık alacak          = brüt prim − tahsil edilen
  Tahsilat oranı %     = Σ tahsil edilen ÷ Σ brüt prim × 100
  Portföy payı %       = grubun brüt primi ÷ toplam brüt prim × 100
  İptal / iade         = eksi tutarlı satırlar; toplamlardan düşer, poliçe adedine girmez
  Uyarı                = şirket payı > --yogunluk (%50) · şirket tahsilat oranı < --tahsilat (%70) · komisyonu yazılmamış poliçe
Komisyon oranları sözleşmenize göredir; kod oran varsaymaz, vergi ve kesinti hesaplamaz.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"no": ("police no", "police numarasi", "police"), "sirket": ("sigorta sirketi", "sirket"), "brans": ("brans", "urun"), "tanzim": ("tanzim tarihi", "duzenleme tarihi", "tarih"), "net": ("net prim",),
         "brut": ("brut prim", "prim", "toplam prim"), "komisyon": ("komisyon", "komisyon tutari"), "oran": ("komisyon orani", "komisyon yuzdesi"), "tahsil": ("tahsil edilen", "tahsilat", "odenen")}
TOPLA = ("brut", "net", "komisyon", "tahsil")


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("sirket", "brut")):
        sirket, no = metin(r["sirket"]), metin(r["no"]) or f"Satır {r['_satir']}"
        if not sirket:
            continue
        d = {a: para(r[a]) if metin(r[a]) else None for a in ("net", "brut", "komisyon", "oran", "tahsil")}
        if d["brut"] is None or any(metin(r[a]) and d[a] is None for a in d) or (d["oran"] is not None and not 0 <= d["oran"] <= 100):
            uy.append(uyari("Yüksek", "Poliçe okunamadı", no, f"Satır {r['_satir']}: tutarlardan biri ya da komisyon oranı okunamadı. Satır alınmadı."))
            continue
        if metin(r["tanzim"]) and tarih(r["tanzim"]) is None:
            uy.append(uyari("Orta", "Tarih okunamadı", no, f"Tanzim tarihi '{metin(r['tanzim'])}' okunamadı; poliçe ay tablosuna girmedi."))
        l.append({"no": no, "sirket": sirket, "brans": metin(r["brans"]) or "(branş yok)", "tanzim": tarih(r["tanzim"]), **d})
    return l, uy


def _oran(pay: Decimal, payda: Decimal) -> Decimal | None:
    return yuvarla(pay * 100 / payda) if payda > 0 else None


def raporla(policeler: list[dict], yogunluk: Decimal = Decimal(50), tahsilat: Decimal = Decimal(70)) -> dict:
    if not policeler:
        raise ValueError("Raporlanacak poliçe yok.")
    for p in policeler:
        p["net"] = p["net"] if p["net"] is not None else p["brut"]
        p["komisyonsuz"] = p["komisyon"] is None and p["oran"] is None
        if p["komisyon"] is None:
            p["komisyon"] = yuvarla(p["net"] * p["oran"] / 100) if p["oran"] is not None else Decimal(0)
        p["tahsil"] = p["tahsil"] if p["tahsil"] is not None else Decimal(0)
        p["alacak"] = p["brut"] - p["tahsil"]
        p["kom_oran"] = _oran(p["komisyon"], p["net"])
    top_brut = sum((p["brut"] for p in policeler), Decimal(0))

    def ozet(l: list[dict]) -> dict:
        t = {a: sum((p[a] for p in l), Decimal(0)) for a in TOPLA}
        return {"police": sum(1 for p in l if p["brut"] > 0), "iptal": sum(1 for p in l if p["brut"] < 0), **t, "kom_oran": _oran(t["komisyon"], t["net"]), "alacak": t["brut"] - t["tahsil"], "tahsilat": _oran(t["tahsil"], t["brut"]),
                "pay": _oran(t["brut"], top_brut), "komisyonsuz": sum(1 for p in l if p["komisyonsuz"])}

    def kirilim(*alanlar: str) -> list[dict]:
        grup = {}
        for p in policeler:
            grup.setdefault(tuple(katla(p[a]) for a in alanlar), []).append(p)
        return sorted(({**{a: g[0][a] for a in alanlar}, **ozet(g)} for g in grup.values()), key=lambda s: (-s["brut"],) + tuple(katla(s[a]) for a in alanlar))

    sirketler, uy = kirilim("sirket"), []
    for s in sirketler:
        if s["pay"] is not None and s["pay"] > yogunluk:
            uy.append(uyari("Bilgi", "Portföy tek şirkette yoğunlaşıyor", s["sirket"], f"Brüt primin %{s['pay']} kadarı bu şirkette (eşik %{yogunluk})."))
        if s["tahsilat"] is not None and s["tahsilat"] < tahsilat:
            uy.append(uyari("Orta", "Tahsilat oranı düşük", s["sirket"], f"Brüt prim {s['brut']:,.2f}, tahsil edilen {s['tahsil']:,.2f} (%{s['tahsilat']}; eşik %{tahsilat}). Açık alacak {s['alacak']:,.2f}."))
        if s["komisyonsuz"]:
            uy.append(uyari("Orta", "Komisyonu yazılmamış poliçe", s["sirket"], f"{s['komisyonsuz']} poliçede komisyon tutarı da oranı da yok; komisyon sıfır sayıldı: "
                            + ", ".join(p["no"] for p in policeler if p["komisyonsuz"] and katla(p["sirket"]) == katla(s["sirket"])) + "."))
    ay = {}
    for p in policeler:
        if p["tanzim"]:
            a = ay.setdefault(f"{p['tanzim']:%Y-%m}", {"brut": Decimal(0), "komisyon": Decimal(0), "police": 0})
            a["brut"], a["komisyon"], a["police"] = a["brut"] + p["brut"], a["komisyon"] + p["komisyon"], a["police"] + (p["brut"] > 0)
    return {"policeler": policeler, "sirketler": sirketler, "branslar": kirilim("brans"), "detay": kirilim("sirket", "brans"), "toplam": ozet(policeler), "aylar": [{"ay": k, **v} for k, v in sorted(ay.items())], "uyarilar": uy}


BASLIK = ["Poliçe", "İptal / İade", "Brüt Prim", "Portföy Payı %", "Net Prim", "Komisyon", "Komisyon %", "Tahsil Edilen", "Açık Alacak", "Tahsilat %"]


def _deger(s: dict) -> list:
    return [s["police"], s["iptal"] or None, s["brut"], s["pay"], s["net"], s["komisyon"], s["kom_oran"], s["tahsil"], s["alacak"], s["tahsilat"]]


def rapor_yaz(cikti: Path, r: dict, tahsilat: Decimal) -> None:
    wb = Workbook()
    for i, (baslik, alan, adlar) in enumerate((("Şirketler", "sirketler", {"sirket": "Sigorta Şirketi"}), ("Branşlar", "branslar", {"brans": "Branş"}), ("Şirket × Branş", "detay", {"sirket": "Sigorta Şirketi", "brans": "Branş"}))):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title, n = baslik, len(adlar)
        _baslik(ws, list(adlar.values()) + BASLIK, (28,) * n + (8, 9, 16, 10, 16, 15, 11, 16, 16, 11), get_column_letter(n + 1) + "2")
        for s in r[alan]:
            _satir(ws, [s[a] for a in adlar] + _deger(s), para_s=range(n + 3, n + 11), dolgu={n + 10: SARI if s["tahsilat"] is not None and s["tahsilat"] < tahsilat else None})
        if n == 1:
            k = _satir(ws, ["TOPLAM"] + _deger(r["toplam"]), para_s=range(4, 12))
            for h in ws[k]:
                h.font = Font(bold=True)
    ws = wb.create_sheet("Aylar")
    _baslik(ws, ["Ay (tanzim)", "Poliçe", "Brüt Prim", "Komisyon"], (14, 9, 18, 16))
    for a in r["aylar"]:
        _satir(ws, [a["ay"], a["police"], a["brut"], a["komisyon"]], para_s=(3, 4))
    ws = wb.create_sheet("Poliçeler")
    _baslik(ws, ["Poliçe No", "Sigorta Şirketi", "Branş", "Tanzim", "Net Prim", "Brüt Prim", "Komisyon", "Komisyon %", "Tahsil Edilen", "Açık Alacak"], (16, 26, 16, 12, 15, 15, 14, 11, 15, 15))
    for p in sorted(r["policeler"], key=lambda p: (katla(p["sirket"]), p["tanzim"] or date.min, p["no"])):
        _satir(ws, [p["no"], p["sirket"], p["brans"], p["tanzim"], p["net"], p["brut"], p["komisyon"], p["kom_oran"], p["tahsil"], p["alacak"]], para_s=(5, 6, 7, 8, 9, 10), tarih_s=(4,),
               dolgu={7: SARI if p["komisyonsuz"] else None, 10: SARI if p["alacak"] > 0 else None})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, yogunluk: Decimal = Decimal(50), tahsilat: Decimal = Decimal(70)) -> dict:
    policeler, uy = oku(yol)
    r = raporla(policeler, yogunluk, tahsilat)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, tahsilat)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Acentenin üretimini sigorta şirketi ve branş bazında toplar; komisyon gelirini ve tahsilat durumunu raporlar.")
    p.add_argument("--policeler", type=Path, help="Poliçe No, Sigorta Şirketi, Branş, Tanzim Tarihi, Net Prim, Brüt Prim, Komisyon, Komisyon Oranı, Tahsil Edilen")
    p.add_argument("--yogunluk", default="50", help="Tek şirketin portföy payı bu yüzdeyi aşarsa bildirilir (varsayılan 50)")
    p.add_argument("--tahsilat", default="70", help="Şirket bazında tahsilat oranı bu yüzdenin altındaysa uyarılır (varsayılan 70)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "acente_portfoy_komisyon_raporu.xlsx")
    a = p.parse_args(argv)
    yol = a.policeler or ORNEK / "policeler.csv"
    yog, tah = para(a.yogunluk), para(a.tahsilat)
    if yog is None or not 0 < yog <= 100 or tah is None or not 0 <= tah <= 100:
        print("[X] --yogunluk ve --tahsilat 0 ile 100 arasında yüzde olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, yog, tah)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {t['police']} poliçe · {len(r['sirketler'])} şirket · brüt prim {t['brut']:,.2f} · komisyon {t['komisyon']:,.2f}" + (f" (%{t['kom_oran']})" if t["kom_oran"] is not None else "") + f" · açık alacak {t['alacak']:,.2f}"
          + (f" · tahsilat %{t['tahsilat']}" if t["tahsilat"] is not None else ""))
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
