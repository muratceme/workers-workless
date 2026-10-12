import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
TARIFE = {("", "depolama"): D(5), ("", "giris"): D(40), ("", "cikis"): D(40)}


def hr(gun, tur, palet, musteri="A", ay=9):
    return {"musteri": musteri, "tarih": date(2026, ay, gun), "tur": tur, "palet": D(palet)}


def ft(hareketler, tarife=TARIFE, hizmetler=(), **kw):
    return main.faturala(hareketler, tarife, list(hizmetler), 2026, 9, **kw)


class FaturaTesti(unittest.TestCase):
    def test_palet_gun(self):
        f = ft([hr(31, "devir", 100, ay=8), hr(10, "giris", 50), hr(20, "cikis", 30)])["faturalar"][0]
        self.assertEqual((f["devir"], f["giris"], f["cikis"], f["kapanis"], f["palet_gun"]), (D(100), D(50), D(30), D(120), D(3720)))                 # 9 × 100 + 10 × 150 + 11 × 120
        self.assertEqual([(s["kalem"].split(" (")[0], s["miktar"], s["tutar"]) for s in f["satirlar"]], [("Depolama", D(3720), D(18600)), ("Giriş elleçleme", D(50), D(2000)), ("Çıkış elleçleme", D(30), D(1200))])
        self.assertEqual((f["ara"], f["kdv"], f["toplam"]), (D(21800), D(4360), D(26160)))
        self.assertEqual(ft([hr(30, "giris", 10)])["faturalar"][0]["palet_gun"], D(10))                                                               # son gün giren 1 gün
        self.assertEqual([ft([hr(1, "giris", 10), hr(1, "cikis", 10)], cikis_dahil=d)["faturalar"][0]["palet_gun"] for d in (False, True)], [D(0), D(10)])
        self.assertEqual(ft([hr(15, "giris", 20, ay=7), hr(20, "cikis", 5, ay=8), hr(5, "giris", 1, ay=10)])["faturalar"][0]["palet_gun"], D(450))    # önceki aylar devir: 15 × 30; sonraki ay yok
        self.assertEqual(ft([hr(1, "devir", 10)])["faturalar"][0]["giris"], D(0))                                                                     # devir elleçleme sayılmaz

    def test_tarife_hizmet_asgari_ve_kdv(self):
        t = TARIFE | {("", "asgari"): D(2000), ("b", "depolama"): D(6)}
        hz = [{"musteri": "A", "hizmet": "Etiket", "adet": D(500), "fiyat": D("1.5"), "tarih": date(2026, 9, 15)}, {"musteri": "A", "hizmet": "Eski", "adet": D(1), "fiyat": D(99), "tarih": date(2026, 8, 1)}]
        r = ft([hr(31, "devir", 100, ay=8), hr(10, "giris", 50), hr(20, "cikis", 30), hr(25, "giris", 10, "B")], t, hz)
        a, b = r["faturalar"]
        self.assertEqual((a["ara"], a["kdv"], a["toplam"], [s["kalem"] for s in a["satirlar"]][-1]), (D(22550), D(4510), D(27060), "Etiket"))
        self.assertEqual(([(s["kalem"].split(" (")[0], s["tutar"]) for s in b["satirlar"]], b["ara"], b["toplam"]), ([("Depolama", D(360)), ("Giriş elleçleme", D(400)), ("Asgari fatura farkı", D(1240))], D(2000), D(2400)))
        self.assertEqual(r["toplam"], {"palet_gun": D(3780), "ara": D(24550), "kdv": D(4910), "toplam": D(29460)})
        self.assertEqual(ft([hr(1, "giris", 1)], kdv=D(0))["faturalar"][0]["kdv"], D(0))
        r = ft([hr(5, "giris", 10), hr(6, "cikis", 15)], {("", "depolama"): D(5)})
        self.assertEqual((sorted(u["tur"] for u in r["uyarilar"]), r["faturalar"][0]["palet_gun"], r["faturalar"][0]["kapanis"]), (["Stok eksiye düştü", "Tarife fiyatı yok", "Tarife fiyatı yok"], D(10), D(-5)))
        with self.assertRaises(ValueError):
            ft([hr(5, "giris", 1, ay=10)])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "d.xlsx"
            r = main.calistir(ORNEK / "stok_hareketleri.csv", ORNEK / "tarife.csv", cikti, 2026, 9, ORNEK / "hizmetler.csv")
            self.assertEqual(([(f["musteri"], f["palet_gun"], f["ara"], f["toplam"]) for f in r["faturalar"]], r["uyarilar"]), ([("Müşteri A", D(3720), D(22550), D(27060)), ("Müşteri B", D(60), D(2000), D(2400))], []))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Fatura Özeti", "Fatura Kalemleri", "Stok Hareketi", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            h, t, z = Path(tmp) / "h.csv", Path(tmp) / "t.csv", Path(tmp) / "z.csv"
            h.write_text("Müşteri;Tarih;Hareket Türü;Palet\nA;01.09.2026;Mal kabul;10\nA;05.09.2026;SEVK;4\nA;dün;Giriş;1\nA;06.09.2026;Sayım;1\nA;01.10.2026;Giriş;5\n", encoding="utf-8")
            t.write_text("Müşteri;Kalem;Fiyat\n;Palet-gün;5\nA;Giriş elleçleme;30\n;Sigorta;9\n;Çıkış;pahalı\n", encoding="utf-8")
            z.write_text("Müşteri;Hizmet;Adet;Birim Fiyat\nA;Shrink;3;20\nA;Bozuk;çok;1\n", encoding="utf-8")
            hl, tf, hz, uy = main.oku(h, t, z)
            self.assertEqual(([(x["tur"], x["palet"]) for x in hl], tf, [(x["hizmet"], x["tarih"]) for x in hz]), ([("giris", D(10)), ("cikis", D(4)), ("giris", D(5))], {("", "depolama"): D(5), ("a", "giris"): D(30)}, [("Shrink", None)]))
            self.assertEqual(sorted(u["tur"] for u in uy), ["Hareket okunamadı", "Hareket okunamadı", "Hizmet satırı okunamadı", "Tarife satırı okunamadı", "Tarife satırı okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            r = main.calistir(h, t, Path(c), 2026, 9, z)
            self.assertEqual((r["faturalar"][0]["palet_gun"], r["faturalar"][0]["ara"], "Dönem sonrası hareketler" in [u["tur"] for u in r["uyarilar"]]), (D(196), D(1340), True))      # 4 × 10 + 26 × 6; 980 + 300 + 0 + 60
            arg = ["--hareketler", str(h), "--tarife", str(t), "--cikti", c]
            self.assertEqual(main.main(arg + ["--donem", "2026-09", "--hizmetler", str(z), "--kdv", "10", "--cikis-gunu-dahil"]), 0)
            self.assertEqual(main.main(arg), 2)                                              # kendi verinizle dönem zorunlu
            self.assertEqual(main.main(arg + ["--donem", "2026-13"]), 2)
            self.assertEqual(main.main(arg + ["--donem", "2025-01"]), 1)
            self.assertEqual(main.main(arg + ["--donem", "2026-09", "--hizmetler", str(Path(tmp) / "yok.csv")]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
