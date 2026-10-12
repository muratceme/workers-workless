"""
Depolama Hizmeti Faturalandırma — Workers / Workless kod bloğu
Depo ve Antrepo Yönetimi › Depo Şefi

Müşteri bazında palet-gün, giriş-çıkış elleçleme ve katma değerli hizmetlerden aylık depolama faturasını hesaplar.
  Gün sonu stoku   = devir + o güne kadarki girişler − çıkışlar           (dönemden önceki hareketler devire eklenir)
  Palet-gün        = dönemin her günü için gün sonu stokunun toplamı       (giriş günü sayılır, çıkış günü sayılmaz; --cikis-gunu-dahil ile çıkış günü de sayılır)
  Depolama bedeli  = palet-gün × palet-gün fiyatı
  Elleçleme        = dönem içi giriş paleti × giriş fiyatı + dönem içi çıkış paleti × çıkış fiyatı
  Hizmetler        = adet × birim fiyat (katma değerli hizmet dosyasından)
  Asgari fatura    = ara toplam asgari tutarın altındaysa fark "Asgari fatura farkı" olarak eklenir
  KDV              = ara toplam × --kdv % (varsayılan 20);   fatura toplamı = ara toplam + KDV
Fiyatlar tarife dosyanızdan gelir: müşteriye özel satır varsa o, yoksa müşterisi boş (genel) satır kullanılır. Kod fatura kesmez; fatura taslağının hesabını çıkarır.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
import calendar  # noqa: E402

HAREKET = {"musteri": ("musteri", "cari", "mal sahibi"), "tarih": ("tarih", "hareket tarihi"), "tur": ("tur", "hareket turu", "islem", "hareket"), "palet": ("palet", "palet adedi", "adet", "miktar")}
HIZMET = {"musteri": HAREKET["musteri"], "hizmet": ("hizmet", "hizmet adi", "aciklama"), "adet": ("adet", "miktar"), "fiyat": ("birim fiyat", "fiyat"), "tarih": ("tarih",)}
TARIFE = {"musteri": HAREKET["musteri"], "kalem": ("kalem", "ucret kalemi", "hizmet"), "fiyat": ("birim fiyat", "fiyat", "tutar")}
KALEM = {"depolama": ("depolama", "palet gun", "palet gunu"), "giris": ("giris", "giris ellecleme"), "cikis": ("cikis", "cikis ellecleme"), "asgari": ("asgari", "asgari fatura", "asgari aylik", "minimum")}
Z = Decimal(0)


def tur_coz(x) -> str | None:
    k = katla(x)
    return "devir" if "devir" in k or "acilis" in k else "giris" if "giris" in k or "kabul" in k else "cikis" if "cikis" in k or "sevk" in k else None


def oku(hareket_yol: Path, tarife_yol: Path, hizmet_yol: Path | None = None) -> tuple[list[dict], dict, list[dict], list[dict]]:
    uy, hareketler, tarife, hizmetler = [], [], {}, []
    for r in kayitlar(hareket_yol, HAREKET, ("musteri", "tarih", "tur", "palet")):
        m = metin(r["musteri"])
        if not m:
            continue
        t, tur, p = tarih(r["tarih"]), tur_coz(r["tur"]), para(r["palet"])
        if t is None or tur is None or p is None or p < 0:
            uy.append(uyari("Yüksek", "Hareket okunamadı", f"{m} ({hareket_yol.name} satır {r['_satir']})", f"Tarih '{metin(r['tarih'])}', tür '{metin(r['tur'])}', palet '{metin(r['palet'])}'. Tür: Giriş, Çıkış veya Devir. Satır alınmadı."))
            continue
        hareketler.append({"musteri": m, "tarih": t, "tur": tur, "palet": p})
    for r in kayitlar(tarife_yol, TARIFE, ("kalem", "fiyat")):
        k, f = katla(r["kalem"]), para(r["fiyat"])
        if not k:
            continue
        kalem = next((a for a, adlar in KALEM.items() if any(k == x or k.startswith(x) for x in adlar)), None)
        if kalem is None or f is None or f < 0:
            uy.append(uyari("Yüksek", "Tarife satırı okunamadı", f"{tarife_yol.name} satır {r['_satir']}", f"Kalem '{metin(r['kalem'])}', fiyat '{metin(r['fiyat'])}'. Kalem: Depolama, Giriş, Çıkış veya Asgari olmalı."))
            continue
        tarife[katla(r["musteri"]), kalem] = f
    if hizmet_yol:
        for r in kayitlar(hizmet_yol, HIZMET, ("musteri", "hizmet", "adet", "fiyat")):
            m, h = metin(r["musteri"]), metin(r["hizmet"])
            if not m or not h:
                continue
            a, f = para(r["adet"]), para(r["fiyat"])
            if a is None or f is None or a < 0 or f < 0:
                uy.append(uyari("Yüksek", "Hizmet satırı okunamadı", f"{m} · {h}", f"Adet '{metin(r['adet'])}', birim fiyat '{metin(r['fiyat'])}'. Satır alınmadı."))
                continue
            hizmetler.append({"musteri": m, "hizmet": h, "adet": a, "fiyat": f, "tarih": tarih(r["tarih"])})
    return hareketler, tarife, hizmetler, uy


def faturala(hareketler: list[dict], tarife: dict, hizmetler: list[dict], yil: int, ay: int, kdv: Decimal = Decimal(20), cikis_dahil: bool = False) -> dict:
    bas, son = date(yil, ay, 1), date(yil, ay, calendar.monthrange(yil, ay)[1])
    mus, uy = {}, []
    for h in hareketler:
        if h["tarih"] <= son:
            mus.setdefault(katla(h["musteri"]), {"musteri": h["musteri"], "h": [], "hizmet": []})["h"].append(h)
    for h in hizmetler:
        if h["tarih"] is None or bas <= h["tarih"] <= son:
            mus.setdefault(katla(h["musteri"]), {"musteri": h["musteri"], "h": [], "hizmet": []})["hizmet"].append(h)
    if not mus:
        raise ValueError(f"{yil}-{ay:02d} dönemi için hareket veya hizmet yok.")
    faturalar = []
    for k, m in mus.items():
        fiyat = {a: tarife.get((k, a), tarife.get(("", a))) for a in KALEM}
        gun = defaultdict(lambda: [Z, Z])
        devir = Z
        for h in m["h"]:
            if h["tarih"] < bas or h["tur"] == "devir":
                devir += -h["palet"] if h["tur"] == "cikis" else h["palet"]
            else:
                gun[h["tarih"]][h["tur"] == "cikis"] += h["palet"]
        stok, palet_gun, giris, cikis, eksi, gunluk = devir, Z, Z, Z, devir < 0, []
        for i in range((son - bas).days + 1):
            g, c = gun.get(bas + timedelta(days=i), (Z, Z))
            stok += g - c
            sayilan = stok + (c if cikis_dahil else Z)
            palet_gun += max(sayilan, Z)
            giris, cikis, eksi = giris + g, cikis + c, eksi or stok < 0
            if g or c or i == 0:
                gunluk.append({"tarih": bas + timedelta(days=i), "giris": g, "cikis": c, "stok": stok})
        if eksi:
            uy.append(uyari("Yüksek", "Stok eksiye düştü", m["musteri"], "Çıkışlar giriş ve devirden fazla. Devir stokunu veya eksik giriş kayıtlarını kontrol edin; eksi stoklu günler palet-gün hesabında 0 sayıldı."))
        satirlar = []
        for ad, miktar, a in (("Depolama (palet-gün)", palet_gun, "depolama"), ("Giriş elleçleme (palet)", giris, "giris"), ("Çıkış elleçleme (palet)", cikis, "cikis")):
            if miktar and fiyat[a] is None:
                uy.append(uyari("Yüksek", "Tarife fiyatı yok", m["musteri"], f"{ad} için fiyat tanımlı değil; bu kalem 0 faturalandı."))
            if miktar:
                satirlar.append({"kalem": ad, "miktar": miktar, "fiyat": fiyat[a] or Z, "tutar": yuvarla(miktar * (fiyat[a] or Z))})
        for h in m["hizmet"]:
            satirlar.append({"kalem": h["hizmet"], "miktar": h["adet"], "fiyat": h["fiyat"], "tutar": yuvarla(h["adet"] * h["fiyat"])})
        ara = sum((s["tutar"] for s in satirlar), Z)
        if fiyat["asgari"] and ara < fiyat["asgari"]:
            satirlar.append({"kalem": "Asgari fatura farkı", "miktar": None, "fiyat": None, "tutar": fiyat["asgari"] - ara})
            ara = fiyat["asgari"]
        vergi = yuvarla(ara * kdv / 100)
        faturalar.append({"musteri": m["musteri"], "devir": devir, "giris": giris, "cikis": cikis, "kapanis": stok, "palet_gun": palet_gun, "satirlar": satirlar, "ara": ara, "kdv": vergi, "toplam": ara + vergi,
                          "gunluk": gunluk})
    faturalar.sort(key=lambda f: f["musteri"])
    return {"faturalar": faturalar, "donem": (bas, son), "kdv_orani": kdv, "cikis_dahil": cikis_dahil, "uyarilar": uy,
            "toplam": {a: sum((f[a] for f in faturalar), Z) for a in ("palet_gun", "ara", "kdv", "toplam")}}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    bas, son = r["donem"]
    ws = wb.active
    ws.title = "Fatura Özeti"
    _baslik(ws, ["Müşteri", "Devir (palet)", "Giriş", "Çıkış", "Dönem Sonu Stok", "Palet-Gün", "Ara Toplam", f"KDV (%{r['kdv_orani'].normalize():f})", "Fatura Toplamı"], (28, 10, 9, 9, 10, 11, 15, 14, 15))
    for f in r["faturalar"]:
        _satir(ws, [f["musteri"], f["devir"], f["giris"], f["cikis"], f["kapanis"], f["palet_gun"], f["ara"], f["kdv"], f["toplam"]], para_s=(7, 8, 9))
    t = r["toplam"]
    _satir(ws, ["Toplam", None, None, None, None, t["palet_gun"], t["ara"], t["kdv"], t["toplam"]], para_s=(7, 8, 9))
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.append([])
    ws.append([f"Dönem {bas:%d.%m.%Y} – {son:%d.%m.%Y}. Palet-gün gün sonu stokuyla hesaplanır; çıkış günü {'sayılır' if r['cikis_dahil'] else 'sayılmaz'}. Fiyatlar tarife dosyanızdan; fatura kesilmeden önce sözleşmeyle karşılaştırın."])
    ws = wb.create_sheet("Fatura Kalemleri")
    _baslik(ws, ["Müşteri", "Kalem", "Miktar", "Birim Fiyat", "Tutar"], (28, 34, 12, 12, 15))
    for f in r["faturalar"]:
        for s in f["satirlar"]:
            _satir(ws, [f["musteri"], s["kalem"], s["miktar"], s["fiyat"], s["tutar"]], para_s=(4, 5), dolgu={2: SARI if s["kalem"] == "Asgari fatura farkı" else None})
        n = _satir(ws, [f["musteri"], "Ara toplam", None, None, f["ara"]], para_s=(5,), dolgu={5: MAVI})
        ws.cell(n, 2).font = Font(bold=True)
    ws = wb.create_sheet("Stok Hareketi")
    _baslik(ws, ["Müşteri", "Tarih", "Giriş", "Çıkış", "Gün Sonu Stok"], (28, 12, 9, 9, 12))
    for f in r["faturalar"]:
        for g in f["gunluk"]:
            _satir(ws, [f["musteri"], g["tarih"], g["giris"] or None, g["cikis"] or None, g["stok"]], tarih_s=(2,), dolgu={5: KIRMIZI if g["stok"] < 0 else None})
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(hareket_yol: Path, tarife_yol: Path, cikti: Path, yil: int, ay: int, hizmet_yol: Path | None = None, **ayar) -> dict:
    hareketler, tarife, hizmetler, uy = oku(hareket_yol, tarife_yol, hizmet_yol)
    r = faturala(hareketler, tarife, hizmetler, yil, ay, **ayar)
    uy += r["uyarilar"]
    sonra = sum(h["tarih"] > r["donem"][1] for h in hareketler)
    if sonra:
        uy.append(uyari("Bilgi", "Dönem sonrası hareketler", hareket_yol.name, f"{sonra} hareket dönemden sonraki tarihli; hesaba girmedi."))
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Müşteri bazında aylık depolama faturası hesabını çıkarır.")
    p.add_argument("--hareketler", type=Path, default=ORNEK / "stok_hareketleri.csv", help="Müşteri, Tarih, Tür (Giriş / Çıkış / Devir), Palet")
    p.add_argument("--tarife", type=Path, default=ORNEK / "tarife.csv", help="Müşteri (boş = genel), Kalem (Depolama / Giriş / Çıkış / Asgari), Birim Fiyat")
    p.add_argument("--hizmetler", type=Path, help="İsteğe bağlı: Müşteri, Hizmet, Adet, Birim Fiyat, Tarih")
    p.add_argument("--donem", help="Fatura dönemi YYYY-AA (örnek veride 2026-09)")
    p.add_argument("--kdv", default="20", help="KDV oranı %% (varsayılan 20)")
    p.add_argument("--cikis-gunu-dahil", action="store_true", help="Çıkış yapılan gün de palet-güne sayılsın")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "depolama_faturasi.xlsx")
    a = p.parse_args(argv)
    ornek = a.hareketler == ORNEK / "stok_hareketleri.csv"
    hizmet = a.hizmetler or (ORNEK / "hizmetler.csv" if ornek else None)
    m = re.fullmatch(r"(\d{4})[-./](\d{1,2})", (a.donem or ("2026-09" if ornek else "")).strip())
    kdv = para(a.kdv)
    if not m or not 1 <= int(m[2]) <= 12 or kdv is None or not 0 <= kdv <= 100:
        print("[X] --donem YYYY-AA biçiminde verilmeli (ör. 2026-09); --kdv 0-100 arasında olmalı")
        return 2
    for y in (a.hareketler, a.tarife, hizmet):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.hareketler, a.tarife, a.cikti, int(m[1]), int(m[2]), hizmet, kdv=kdv, cikis_dahil=a.cikis_gunu_dahil)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    for f in r["faturalar"]:
        print(f"[OK] {f['musteri']}: {f['palet_gun']:f} palet-gün · ara toplam {f['ara']:,.2f} · KDV {f['kdv']:,.2f} · toplam {f['toplam']:,.2f}")
    print(f"[OK] {len(r['faturalar'])} müşteri · fatura toplamı {r['toplam']['toplam']:,.2f} · rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
