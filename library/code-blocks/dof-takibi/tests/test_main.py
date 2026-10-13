import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)


def g(n):
    return None if n is None else BUGUN + timedelta(days=n)


def d(no, termin=30, tamam=None, dogr=None, sonuc=None, acilis=-60, sorumlu="S", kaynak="K", tur="Düzeltici"):
    """Tarihler: bugünden kaç gün sonra (eksi = geçmiş; None = boş)."""
    return {"no": no, "tur": tur, "kaynak": kaynak, "konu": "Konu", "sorumlu": sorumlu, "acilis": g(acilis), "termin": g(termin), "tamam": g(tamam), "dogr": g(dogr), "sonuc": sonuc}


def veri():
    return [d("d1", -31), d("d2", -1), d("d3", 15), d("d4", 16), d("d5", None), d("d6", -35, -30), d("d7", -20, -31), d("d8", -40, -40, -10, "Etkin"), d("d9", -50, -45, -10, "Etkisiz"), d("d10", -30, -35, -5)]


class DegerlendirTesti(unittest.TestCase):
    def test_durumlar(self):
        r = main.degerlendir(veri(), BUGUN)
        self.assertEqual({x["no"]: (x["durum"], x["gecikme"], x["bekleme"]) for x in r["doflar"]},
                         {"d1": ("Gecikmiş", 31, None), "d2": ("Gecikmiş", 1, None), "d3": ("Termin yaklaşıyor", None, None), "d4": ("Açık", None, None), "d5": ("Termin yok", None, None), "d6": ("Doğrulama bekliyor", None, 30),
                          "d7": ("Doğrulama gecikti", None, 31), "d8": ("Kapalı (etkin)", None, None), "d9": ("Etkisiz, yeniden açılmalı", None, None), "d10": ("Sonuç yazılmamış", None, None)})
        self.assertEqual([x["no"] for x in r["doflar"]], ["d9", "d1", "d2", "d7", "d10", "d5", "d3", "d6", "d4", "d8"])
        self.assertEqual(r["ozet"], {"dof": 10, "acik": 9, "gecikmis": 2, "dogrulama": 2, "kapali": 1, "etkisiz": 1, "kapanma": D("10.00"), "etkinlik": D("50.00"), "zamaninda": D("60.00"), "ort_aksiyon": D("23.80"), "ort_kapanma": D("50.00")})
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         sorted([("Yüksek", "Gecikmiş faaliyet", "d1 · S"), ("Orta", "Gecikmiş faaliyet", "d2 · S"), ("Bilgi", "Termin yaklaşıyor", "d3 · S"), ("Orta", "Termin yok", "d5 · S"), ("Orta", "Etkinlik doğrulaması gecikti", "d7 · S"),
                                 ("Yüksek", "Etkisiz bulunan faaliyet", "d9 · S"), ("Orta", "Etkinlik sonucu yazılmamış", "d10 · S")]))
        self.assertEqual([(x["durum"], x["dof"]) for x in r["durumlar"]],
                         [("Etkisiz, yeniden açılmalı", 1), ("Gecikmiş", 2), ("Doğrulama gecikti", 1), ("Sonuç yazılmamış", 1), ("Termin yok", 1), ("Termin yaklaşıyor", 1), ("Doğrulama bekliyor", 1), ("Açık", 1), ("Kapalı (etkin)", 1)])

    def test_esikler_ve_kirilim(self):
        s = lambda **kw: {x["no"]: x["durum"] for x in main.degerlendir([d("a", 16), d("b", -40, -30)], BUGUN, **kw)["doflar"]}   # noqa: E731
        self.assertEqual((s(), s(gun=16), s(dogrulama=29)), ({"a": "Açık", "b": "Doğrulama bekliyor"}, {"a": "Termin yaklaşıyor", "b": "Doğrulama bekliyor"}, {"a": "Açık", "b": "Doğrulama gecikti"}))
        o = lambda **kw: [u["onem"] for u in main.degerlendir([d("a", -29)], BUGUN, **kw)["uyarilar"]]   # noqa: E731
        self.assertEqual((o(), o(kritik=29)), (["Orta"], ["Yüksek"]))
        r = main.degerlendir([d("a", -10, -70, -75, "Etkin"), d("b", -10, None, -5, "Etkin")], BUGUN)
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]), [("Tarih sırası tutarsız", "a · S")] * 2 + [("Tarih sırası tutarsız", "b · S")])
        l = [d("1", -5, sorumlu="A", kaynak="Şikâyet"), d("2", -40, -40, -10, "Etkin", sorumlu="a", kaynak="Denetim", tur="Önleyici"), d("3", 50, sorumlu="B", kaynak="şikâyet"), d("4", -40, -30, -10, "Etkisiz", sorumlu="B", kaynak="Denetim")]
        r = main.degerlendir(l, BUGUN)
        k = lambda alan: [(x["ad"], x["dof"], x["acik"], x["gecikmis"], x["kapali"], x["etkisiz"], x["kapanma"], x["etkinlik"], x["zamaninda"]) for x in r[alan]]   # noqa: E731
        self.assertEqual(k("sorumlular"), [("A", 2, 1, 1, 1, 0, D("50.00"), D("100.00"), D("100.00")), ("B", 2, 2, 0, 0, 1, D("0.00"), D("0.00"), D("0.00"))])
        self.assertEqual(k("kaynaklar"), [("Şikâyet", 2, 2, 1, 0, 0, D("0.00"), None, None), ("Denetim", 2, 1, 0, 1, 1, D("50.00"), D("50.00"), D("50.00"))])
        self.assertEqual([(x["ad"], x["dof"]) for x in r["turler"]], [("Düzeltici", 3), ("Önleyici", 1)])
        with self.assertRaises(ValueError):
            main.degerlendir([], BUGUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "dof_listesi.csv", cikti, BUGUN)
            self.assertEqual([(x["no"][-3:], x["durum"], x["gecikme"], x["bekleme"]) for x in r["doflar"]],
                             [("035", "Etkisiz, yeniden açılmalı", None, None), ("031", "Gecikmiş", 32, None), ("038", "Gecikmiş", 7, None), ("037", "Doğrulama gecikti", None, 32), ("036", "Termin yok", None, None),
                              ("033", "Termin yaklaşıyor", None, None), ("034", "Doğrulama bekliyor", None, 18), ("039", "Açık", None, None), ("032", "Kapalı (etkin)", None, None), ("040", "Kapalı (etkin)", None, None)])
            o = r["ozet"]
            self.assertEqual((o["dof"], o["acik"], o["gecikmis"], o["dogrulama"], o["kapali"], o["etkisiz"], o["kapanma"], o["etkinlik"], o["zamaninda"]), (10, 8, 2, 2, 2, 1, D("20.00"), D("66.67"), D("80.00")))
            self.assertEqual([(x["ad"], x["dof"], x["acik"], x["gecikmis"], x["etkisiz"]) for x in r["sorumlular"]],
                             [("Üretim Müdürü", 3, 3, 2, 0), ("Kalite Mühendisi", 3, 2, 0, 1), ("Bakım Mühendisi", 2, 1, 0, 0), ("İK Uzmanı", 1, 1, 0, 0), ("Satın Alma Uzmanı", 1, 1, 0, 0)])
            self.assertEqual([(u["onem"], u["tur"], u["kim"][7:10]) for u in r["uyarilar"]],
                             [("Yüksek", "Etkisiz bulunan faaliyet", "035"), ("Yüksek", "Gecikmiş faaliyet", "031"), ("Orta", "Etkinlik doğrulaması gecikti", "037"), ("Orta", "Gecikmiş faaliyet", "038"), ("Orta", "Termin yok", "036"),
                              ("Bilgi", "Termin yaklaşıyor", "033")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["DÖF Listesi", "Özet", "Sorumlular", "Kaynaklar", "Türler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "d.csv"
            y.write_text("Faaliyet No;DÖF Türü;Uygunsuzluk Kaynağı;Tanım;Sorumlu Kişi;DÖF Tarihi;Hedef Tarih;Gerçekleşme Tarihi;Doğrulama Tarihi;Doğrulama Sonucu\nD1;Düzeltici;Şikâyet;Metin;Ali;01.09.2026;x;10.09.2026;20.09.2026;ETKİLİ\n"
                         "D2;;;;;02.09.2026;;;;Belki\nD3;;;;;02.09.2026;;;;Etkin değil\nd1;;;;;03.09.2026;;;;\nD4;;;;;x;;;;\n;;;;;03.09.2026;;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["tur"], x["kaynak"], x["konu"], x["sorumlu"], x["acilis"], x["termin"], x["tamam"], x["dogr"], x["sonuc"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("D1", "Düzeltici", "Şikâyet", "Metin", "Ali", date(2026, 9, 1), None, date(2026, 9, 10), date(2026, 9, 20), "Etkin"), ("D2", main.YOK, main.YOK, "", main.YOK, date(2026, 9, 2), None, None, None, None),
                               ("D3", main.YOK, main.YOK, "", main.YOK, date(2026, 9, 2), None, None, None, "Etkisiz")],
                              [("Tarih okunamadı", "D1"), ("Etkinlik sonucu tanınmadı", "D2"), ("DÖF okunamadı", "d1"), ("DÖF okunamadı", "D4")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--doflar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "12.10.2026", "--gun", "7", "--dogrulama", "60", "--kritik", "15"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--gun", "-1"], ["--dogrulama", "-1"], ["--kritik", "0"])], [2] * 4)
            self.assertEqual(main.main(["--doflar", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("DÖF No;Konu\nA;B\n", "DÖF No;Açılış Tarihi\n;01.03.2026\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
