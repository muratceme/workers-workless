"""
DÖF Takibi — Gerisi Bende kod bloğu
Kalite Güvence ve Kalite Kontrol › Kalite Mühendisi

Düzeltici ve önleyici faaliyetleri sorumlu, termin ve etkinlik doğrulamasıyla izler.
  Aksiyon tamamlanmamış  : Gecikmiş (termin < durum tarihi) · Termin yok · Termin yaklaşıyor (termine --gun, 15 gün veya az) · Açık
  Aksiyon tamamlanmış    : Doğrulama bekliyor · Doğrulama gecikti (tamamlanmanın üzerinden --dogrulama, 30 günden fazla geçmiş)
  Doğrulama yapılmış     : Kapalı (etkin) · Etkisiz, yeniden açılmalı · Sonuç yazılmamış
  Kapanma oranı %        = kapalı (etkin) ÷ toplam × 100
  Etkinlik oranı %       = etkin bulunan ÷ doğrulaması sonuçlanan × 100
  Zamanında aksiyon %    = terminde ya da önce tamamlanan ÷ termini olan tamamlanmış faaliyet × 100
  Ortalama aksiyon süresi = aksiyon tamamlanma − açılış;   ortalama kapanma süresi = etkinlik doğrulama − açılış (kapalı faaliyetler)
Faaliyetin etkin olup olmadığına siz karar verirsiniz; kod kök neden analizini ve kanıtı değerlendirmez.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"no": ("dof no", "dof numarasi", "faaliyet no", "no"), "tur": ("tur", "faaliyet turu", "dof turu"), "kaynak": ("kaynak", "dof kaynagi", "uygunsuzluk kaynagi"), "konu": ("konu", "uygunsuzluk", "aciklama", "tanim"),
         "sorumlu": ("sorumlu", "faaliyet sorumlusu", "sorumlu kisi"), "acilis": ("acilis tarihi", "dof tarihi", "tarih"), "termin": ("termin", "termin tarihi", "hedef tarih"),
         "tamam": ("aksiyon tamamlanma tarihi", "tamamlanma tarihi", "gerceklesme tarihi"), "dogr": ("etkinlik dogrulama tarihi", "dogrulama tarihi", "etkinlik kontrol tarihi"),
         "sonuc": ("etkinlik sonucu", "dogrulama sonucu", "etkinlik")}
DURUM = ("Etkisiz, yeniden açılmalı", "Gecikmiş", "Doğrulama gecikti", "Sonuç yazılmamış", "Termin yok", "Termin yaklaşıyor", "Doğrulama bekliyor", "Açık", "Kapalı (etkin)")
DURUM_RENK = {"Etkisiz, yeniden açılmalı": KIRMIZI, "Gecikmiş": KIRMIZI, "Doğrulama gecikti": SARI, "Sonuç yazılmamış": SARI, "Termin yok": SARI, "Termin yaklaşıyor": SARI, "Doğrulama bekliyor": MAVI, "Açık": None,
              "Kapalı (etkin)": YESIL}
SONUC = {"etkin": "Etkin", "etkili": "Etkin", "uygun": "Etkin", "evet": "Etkin", "etkisiz": "Etkisiz", "etkili degil": "Etkisiz", "etkin degil": "Etkisiz", "uygun degil": "Etkisiz", "hayir": "Etkisiz"}
YOK = "(yazılmamış)"


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("no", "acilis")):
        no = metin(r["no"])
        if not no:
            continue
        acilis = tarih(r["acilis"])
        if acilis is None or katla(no) in gorulen:
            uy.append(uyari("Yüksek", "DÖF okunamadı", no, f"Satır {r['_satir']}: " + ("aynı DÖF numarası daha önce geçti" if katla(no) in gorulen else f"açılış tarihi '{metin(r['acilis'])}' okunamadı") + ". Satır alınmadı."))
            continue
        gorulen.add(katla(no))
        for alan, ad in (("termin", "Termin"), ("tamam", "Aksiyon tamamlanma tarihi"), ("dogr", "Etkinlik doğrulama tarihi")):
            if metin(r[alan]) and tarih(r[alan]) is None:
                uy.append(uyari("Orta", "Tarih okunamadı", no, f"{ad} '{metin(r[alan])}' okunamadı; boş sayıldı."))
        sonuc = SONUC.get(katla(r["sonuc"]))
        if sonuc is None and metin(r["sonuc"]):
            uy.append(uyari("Orta", "Etkinlik sonucu tanınmadı", no, f"'{metin(r['sonuc'])}' Etkin ya da Etkisiz olarak okunamadı; sonuç yazılmamış sayıldı."))
        l.append({"no": no, "tur": metin(r["tur"]) or YOK, "kaynak": metin(r["kaynak"]) or YOK, "konu": metin(r["konu"]), "sorumlu": metin(r["sorumlu"]) or YOK, "acilis": acilis, "termin": tarih(r["termin"]), "tamam": tarih(r["tamam"]),
                  "dogr": tarih(r["dogr"]), "sonuc": sonuc})
    return l, uy


def _ort(d: list[int]) -> Decimal | None:
    return yuvarla(Decimal(sum(d)) / len(d)) if d else None


def _oran(pay: int, payda: int) -> Decimal | None:
    return yuvarla(Decimal(pay * 100) / payda) if payda else None


def degerlendir(doflar: list[dict], bugun: date, gun: int = 15, dogrulama: int = 30, kritik: int = 30) -> dict:
    if not doflar:
        raise ValueError("İzlenecek DÖF yok.")
    uy = []
    for d in doflar:
        kim = f"{d['no']} · {d['sorumlu']}"
        d["gecikme"] = d["bekleme"] = None
        d["yas"] = max((bugun - d["acilis"]).days, 0)
        if d["dogr"] is not None:
            d["durum"] = "Kapalı (etkin)" if d["sonuc"] == "Etkin" else "Etkisiz, yeniden açılmalı" if d["sonuc"] == "Etkisiz" else "Sonuç yazılmamış"
            if d["sonuc"] == "Etkisiz":
                uy.append(uyari("Yüksek", "Etkisiz bulunan faaliyet", kim, f"{d['dogr']:%d.%m.%Y} tarihli doğrulamada etkisiz bulundu. Kök nedeni yeniden değerlendirip yeni faaliyet açın. {d['konu']}"))
            elif d["sonuc"] is None:
                uy.append(uyari("Orta", "Etkinlik sonucu yazılmamış", kim, f"Doğrulama tarihi {d['dogr']:%d.%m.%Y} ama sonuç (Etkin / Etkisiz) yok. {d['konu']}"))
        elif d["tamam"] is not None:
            d["bekleme"] = max((bugun - d["tamam"]).days, 0)
            d["durum"] = "Doğrulama gecikti" if d["bekleme"] > dogrulama else "Doğrulama bekliyor"
            if d["durum"] == "Doğrulama gecikti":
                uy.append(uyari("Orta", "Etkinlik doğrulaması gecikti", kim, f"Aksiyon {d['tamam']:%d.%m.%Y} tarihinde tamamlandı; {d['bekleme']} gündür etkinlik doğrulaması yapılmadı. {d['konu']}"))
        else:
            kalan = (d["termin"] - bugun).days if d["termin"] else None
            d["durum"] = "Termin yok" if kalan is None else "Gecikmiş" if kalan < 0 else "Termin yaklaşıyor" if kalan <= gun else "Açık"
            if d["durum"] == "Gecikmiş":
                d["gecikme"] = -kalan
                uy.append(uyari("Yüksek" if d["gecikme"] >= kritik else "Orta", "Gecikmiş faaliyet", kim, f"Termin {d['termin']:%d.%m.%Y} ({d['gecikme']} gün geçti); DÖF {d['yas']} gündür açık. {d['konu']}"))
            elif d["durum"] == "Termin yok":
                uy.append(uyari("Orta", "Termin yok", kim, f"DÖF {d['yas']} gündür açık ama termin verilmemiş. {d['konu']}"))
            elif d["durum"] == "Termin yaklaşıyor":
                uy.append(uyari("Bilgi", "Termin yaklaşıyor", kim, f"Termin {d['termin']:%d.%m.%Y} ({kalan} gün kaldı). {d['konu']}"))
        for once, sonra, ad in (("acilis", "tamam", "Aksiyon tamamlanma tarihi açılıştan"), ("tamam", "dogr", "Doğrulama tarihi aksiyon tamamlanmadan")):
            if d[once] and d[sonra] and d[sonra] < d[once]:
                uy.append(uyari("Orta", "Tarih sırası tutarsız", kim, f"{ad} önce ({d[sonra]:%d.%m.%Y} < {d[once]:%d.%m.%Y}). Tarihleri kontrol edin."))
        if d["dogr"] and not d["tamam"]:
            uy.append(uyari("Orta", "Tarih sırası tutarsız", kim, "Etkinlik doğrulama tarihi var ama aksiyon tamamlanma tarihi yok."))

    def ozet(l: list[dict]) -> dict:
        kapali, etkisiz = [d for d in l if d["durum"] == "Kapalı (etkin)"], sum(1 for d in l if d["durum"] == "Etkisiz, yeniden açılmalı")
        terminli = [d for d in l if d["tamam"] and d["termin"]]
        return {"dof": len(l), "acik": len(l) - len(kapali), "gecikmis": sum(1 for d in l if d["durum"] == "Gecikmiş"), "dogrulama": sum(1 for d in l if d["durum"] in ("Doğrulama bekliyor", "Doğrulama gecikti")),
                "kapali": len(kapali), "etkisiz": etkisiz, "kapanma": _oran(len(kapali), len(l)), "etkinlik": _oran(len(kapali), len(kapali) + etkisiz), "zamaninda": _oran(sum(1 for d in terminli if d["tamam"] <= d["termin"]), len(terminli)),
                "ort_aksiyon": _ort([max((d["tamam"] - d["acilis"]).days, 0) for d in l if d["tamam"]]), "ort_kapanma": _ort([max((d["dogr"] - d["acilis"]).days, 0) for d in kapali])}

    def kirilim(alan: str) -> list[dict]:
        grup = {}
        for d in doflar:
            grup.setdefault(katla(d[alan]), []).append(d)
        return sorted(({"ad": g[0][alan], **ozet(g)} for g in grup.values()), key=lambda s: (-s["gecikmis"], -s["acik"], -s["dof"], katla(s["ad"])))

    return {"doflar": sorted(doflar, key=lambda d: (DURUM.index(d["durum"]), -(d["gecikme"] or 0), -(d["bekleme"] or 0), katla(d["no"]))), "ozet": ozet(doflar), "sorumlular": kirilim("sorumlu"), "kaynaklar": kirilim("kaynak"),
            "turler": kirilim("tur"), "durumlar": [{"durum": x, "dof": sum(1 for d in doflar if d["durum"] == x)} for x in DURUM], "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, bugun: date) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "DÖF Listesi"
    _baslik(ws, ["Durum", "DÖF No", "Tür", "Kaynak", "Konu", "Sorumlu", "Açılış", "Termin", "Aksiyon Tamamlanma", "Etkinlik Doğrulama", "Etkinlik Sonucu", "Gecikme (gün)", "Doğrulama Bekleme (gün)", "Not (siz doldurun)"],
            (26, 13, 11, 20, 44, 22, 12, 12, 12, 12, 10, 10, 11, 34))
    for d in r["doflar"]:
        _satir(ws, [d["durum"], d["no"], d["tur"], d["kaynak"], d["konu"], d["sorumlu"], d["acilis"], d["termin"], d["tamam"], d["dogr"], d["sonuc"], d["gecikme"], d["bekleme"], None], tarih_s=(7, 8, 9, 10), dolgu={1: DURUM_RENK[d["durum"]]}, sar=(5,))
    ws = wb.create_sheet("Özet")
    o = r["ozet"]
    _baslik(ws, ["Gösterge", "Değer"], (44, 14))
    for etiket, deger, para_mi in (("Toplam DÖF", o["dof"], False), ("Açık DÖF (kapalı olmayan)", o["acik"], False), ("Gecikmiş faaliyet", o["gecikmis"], False), ("Etkinlik doğrulaması bekleyen", o["dogrulama"], False),
                                   ("Kapalı (etkin)", o["kapali"], False), ("Etkisiz bulunan", o["etkisiz"], False), ("Kapanma oranı %", o["kapanma"], True), ("Etkinlik oranı %", o["etkinlik"], True),
                                   ("Zamanında tamamlanan aksiyon %", o["zamaninda"], True), ("Ortalama aksiyon süresi (gün)", o["ort_aksiyon"], True), ("Ortalama kapanma süresi (gün)", o["ort_kapanma"], True)):
        _satir(ws, [etiket, deger], para_s=(2,) if para_mi else ())
    ws.append([])
    _baslik(ws, ["Durum", "DÖF"], dondur="A2")
    for d in r["durumlar"]:
        _satir(ws, [d["durum"], d["dof"]], dolgu={1: DURUM_RENK[d["durum"]] if d["dof"] else None})
    ws.append([])
    ws.append([f"Durum tarihi: {bugun:%d.%m.%Y}"])
    for baslik, alan, ad in (("Sorumlular", "sorumlular", "Sorumlu"), ("Kaynaklar", "kaynaklar", "Kaynak"), ("Türler", "turler", "Tür")):
        ws = wb.create_sheet(baslik)
        _baslik(ws, [ad, "DÖF", "Açık", "Gecikmiş", "Doğrulama Bekleyen", "Kapalı (etkin)", "Etkisiz", "Kapanma %", "Etkinlik %", "Zamanında Aksiyon %", "Ort. Aksiyon Süresi (gün)", "Ort. Kapanma Süresi (gün)"],
                (28, 7, 7, 9, 11, 9, 8, 10, 10, 11, 12, 12))
        for s in r[alan]:
            _satir(ws, [s["ad"], s["dof"], s["acik"] or None, s["gecikmis"] or None, s["dogrulama"] or None, s["kapali"] or None, s["etkisiz"] or None, s["kapanma"], s["etkinlik"], s["zamaninda"], s["ort_aksiyon"], s["ort_kapanma"]],
                   para_s=(8, 9, 10, 11, 12), dolgu={4: KIRMIZI if s["gecikmis"] else None, 7: KIRMIZI if s["etkisiz"] else None})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, gun: int = 15, dogrulama: int = 30, kritik: int = 30) -> dict:
    doflar, uy = oku(yol)
    r = degerlendir(doflar, bugun, gun, dogrulama, kritik)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, bugun)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Düzeltici ve önleyici faaliyetleri sorumlu, termin ve etkinlik doğrulamasıyla izler.")
    p.add_argument("--doflar", type=Path, help="DÖF No, Tür, Kaynak, Konu, Sorumlu, Açılış Tarihi, Termin, Aksiyon Tamamlanma Tarihi, Etkinlik Doğrulama Tarihi, Etkinlik Sonucu")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--gun", type=int, default=15, help="Termine kaç gün kala 'yaklaşıyor' sayılsın (varsayılan 15)")
    p.add_argument("--dogrulama", type=int, default=30, help="Aksiyon tamamlandıktan sonra kaç günden fazla doğrulamasız kalırsa gecikmiş sayılsın (varsayılan 30)")
    p.add_argument("--kritik", type=int, default=30, help="Bu kadar gün ve üstü termin gecikmesi yüksek önemli sayılır (varsayılan 30)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "dof_takibi.xlsx")
    a = p.parse_args(argv)
    ornek = a.doflar is None
    yol = ORNEK / "dof_listesi.csv" if ornek else a.doflar
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    if bugun is None or a.gun < 0 or a.dogrulama < 0 or a.kritik < 1:
        print("[X] --tarih GG.AA.YYYY; --gun ve --dogrulama eksi olamaz; --kritik en az 1 olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, bugun, a.gun, a.dogrulama, a.kritik)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {o['dof']} DÖF · açık {o['acik']} (gecikmiş {o['gecikmis']}, doğrulama bekleyen {o['dogrulama']}) · kapalı {o['kapali']} · etkisiz {o['etkisiz']}" + (f" · etkinlik %{o['etkinlik']}" if o["etkinlik"] is not None else ""))
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
