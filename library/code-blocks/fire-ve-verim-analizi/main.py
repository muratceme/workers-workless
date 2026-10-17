"""
Fire ve Verim Analizi — Gerisi Bende kod bloğu
Üretim › Üretim Mühendisi

Hammadde girişi ile mamul çıkışından ürün ve hat bazında fire oranını ve verim kaybı maliyetini hesaplar.
  Fire (kg)            = hammadde girişi − mamul çıkışı − geri kazanılan
  Fire %               = fire ÷ hammadde girişi × 100          Verim % = mamul çıkışı ÷ hammadde girişi × 100
  Fire maliyeti        = fire × hammadde birim maliyeti
  Standart üstü fire   = fire − hammadde girişi × standart fire %      (kayıt bazında; eksiyse 0)
  Verim kaybı maliyeti = standart üstü fire × hammadde birim maliyeti
Gruplarda oranlar toplamlardan hesaplanır (ağırlıklı); standart üstü fire kayıt kayıt toplanır, iyi günler kötü günleri silmez.
Maliyet yalnız hammaddedir; işçilik, enerji ve makine zamanı kaybı dahil değildir.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"tarih": ("tarih", "donem"), "hat": ("hat", "makine", "is merkezi"), "urun": ("urun", "urun adi", "stok kodu"),
         "giris": ("hammadde girisi", "hammadde girisi kg", "giris", "giris kg", "kullanilan hammadde"), "cikis": ("mamul cikisi", "mamul cikisi kg", "cikis", "cikis kg", "saglam uretim"),
         "geri": ("geri kazanilan", "geri kazanilan kg", "geri donusum", "kirma"), "maliyet": ("birim maliyet", "birim maliyet tl kg", "hammadde birim maliyeti", "maliyet"),
         "standart": ("standart fire", "standart fire yuzde", "hedef fire", "standart")}
Z = Decimal(0)


def oku(yol: Path, standart: Decimal | None = None) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("urun", "giris", "cikis")):
        if not metin(r["urun"]):
            continue
        kim = f"{metin(r['urun'])} (satır {r['_satir']})"
        g, c, geri, m, s = (para(r[a]) for a in ("giris", "cikis", "geri", "maliyet", "standart"))
        if g is None or c is None or g <= 0 or c < 0:
            uy.append(uyari("Yüksek", "Miktar okunamadı", kim, f"Giriş '{metin(r['giris'])}', çıkış '{metin(r['cikis'])}'. Giriş pozitif, çıkış sıfır veya pozitif olmalı. Kayıt alınmadı."))
            continue
        geri = geri or Z
        if g - c - geri < 0:
            uy.append(uyari("Yüksek", "Çıkış girişten büyük", kim, f"Giriş {g:f}, çıkış {c:f}, geri kazanılan {geri:f}: fire eksi çıkıyor. Birim veya dönem kayması olabilir; kayıt alınmadı."))
            continue
        if m is None and metin(r["maliyet"]):
            uy.append(uyari("Orta", "Birim maliyet okunamadı", kim, f"'{metin(r['maliyet'])}'. Maliyet hesaplanmadı."))
        if metin(r["standart"]) and (s is None or not 0 <= s < 100):
            uy.append(uyari("Orta", "Standart fire okunamadı", kim, f"'{metin(r['standart'])}'. Varsayılan kullanıldı."))
            s = None
        l.append({"tarih": tarih(r["tarih"]), "hat": metin(r["hat"]) or "Belirtilmemiş", "urun": metin(r["urun"]), "giris": g, "cikis": c, "geri": geri, "maliyet": m,
                  "standart": s if s is not None else standart})
    return l, uy


def hesapla(k: dict) -> dict:
    k["fire"] = k["giris"] - k["cikis"] - k["geri"]
    k["ustu"] = max(k["fire"] - k["giris"] * k["standart"] / 100, Z) if k["standart"] is not None else None
    k["fire_tl"] = yuvarla(k["fire"] * k["maliyet"]) if k["maliyet"] is not None else None
    k["ustu_tl"] = yuvarla(k["ustu"] * k["maliyet"]) if k["maliyet"] is not None and k["ustu"] is not None else None
    return k | _oran(k)


def _oran(x: dict) -> dict:
    return {"fire_oran": yuvarla(x["fire"] / x["giris"] * 100), "verim": yuvarla(x["cikis"] / x["giris"] * 100)}


def topla(kayit: list[dict], ad="Toplam") -> dict:
    x = {"ad": ad, "kayit": len(kayit)} | {a: sum((k[a] for k in kayit), Z) for a in ("giris", "cikis", "geri", "fire")}
    for a in ("ustu", "fire_tl", "ustu_tl"):
        v = [k[a] for k in kayit if k[a] is not None]
        x[a] = sum(v, Z) if v else None
    x["asan"] = sum(1 for k in kayit if k["ustu"])
    return x | _oran(x)


def analiz(kayit: list[dict]) -> dict:
    if not kayit:
        raise ValueError("Analiz edilecek kayıt yok.")
    kayit = [hesapla(k) for k in kayit]

    def grup(f):
        g = defaultdict(list)
        for k in kayit:
            g[f(k)].append(k)
        return sorted((topla(v, a) for a, v in g.items()), key=lambda x: (-(x["ustu_tl"] or Z), -x["fire_oran"], str(x["ad"])))

    return {"kayit": kayit, "toplam": topla(kayit), "urun": grup(lambda k: k["urun"]), "hat": grup(lambda k: k["hat"]), "urun_hat": grup(lambda k: f"{k['urun']} · {k['hat']}")}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb, t = Workbook(), r["toplam"]
    ws = wb.active
    ws.title = "Özet"
    _baslik(ws, ["Gösterge", "Değer"], (40, 20))
    for ad, v in (("Kayıt sayısı", t["kayit"]), ("Hammadde girişi (kg)", t["giris"]), ("Mamul çıkışı (kg)", t["cikis"]), ("Geri kazanılan (kg)", t["geri"]), ("Fire (kg)", t["fire"]),
                  ("Fire %", t["fire_oran"]), ("Verim %", t["verim"]), ("Fire maliyeti (TL)", t["fire_tl"]), ("Standart üstü fire (kg)", t["ustu"]),
                  ("Verim kaybı maliyeti (TL, standart üstü)", t["ustu_tl"]), ("Standardı aşan kayıt", t["asan"])):
        _satir(ws, [ad, v], para_s=(2,) if isinstance(v, Decimal) else ())
    ws.append([])
    ws.append(["Maliyet yalnız hammaddedir. Standart üstü fire kayıt bazında toplanır; standardın altında kalan kayıtlar aşanları dengelemez."])
    bas = ["Kayıt", "Hammadde Girişi (kg)", "Mamul Çıkışı (kg)", "Geri Kazanılan (kg)", "Fire (kg)", "Fire %", "Verim %", "Fire Maliyeti (TL)", "Standart Üstü Fire (kg)", "Verim Kaybı Maliyeti (TL)",
           "Standardı Aşan Kayıt"]
    for sayfa, veri, kol in (("Ürün Bazında", r["urun"], "Ürün"), ("Hat Bazında", r["hat"], "Hat"), ("Ürün-Hat", r["urun_hat"], "Ürün · Hat")):
        ws = wb.create_sheet(sayfa)
        _baslik(ws, [kol, *bas], (30, 7, 16, 16, 14, 12, 9, 9, 16, 14, 18, 10))
        for x in veri:
            _satir(ws, [x["ad"], x["kayit"], x["giris"], x["cikis"], x["geri"], x["fire"], x["fire_oran"], x["verim"], x["fire_tl"], x["ustu"], x["ustu_tl"], x["asan"] or None],
                   para_s=range(3, 12), dolgu={11: KIRMIZI if x["ustu_tl"] else None})
    ws = wb.create_sheet("Kayıtlar")
    _baslik(ws, ["Tarih", "Hat", "Ürün", "Giriş (kg)", "Çıkış (kg)", "Geri Kazanılan (kg)", "Fire (kg)", "Fire %", "Standart Fire %", "Standart Üstü (kg)", "Birim Maliyet", "Fire Maliyeti",
                 "Verim Kaybı Maliyeti", "Neden (siz doldurun)"], (12, 14, 22, 12, 12, 12, 12, 9, 10, 12, 12, 14, 14, 30))
    for k in sorted(r["kayit"], key=lambda k: -(k["ustu_tl"] or k["ustu"] or Z)):
        _satir(ws, [k["tarih"], k["hat"], k["urun"], k["giris"], k["cikis"], k["geri"], k["fire"], k["fire_oran"], k["standart"], k["ustu"], k["maliyet"], k["fire_tl"], k["ustu_tl"], ""],
               para_s=range(4, 14), tarih_s=(1,), dolgu={8: KIRMIZI if k["ustu"] else None})
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, standart: Decimal | None = None) -> dict:
    kayit, uy = oku(yol, standart)
    r = analiz(kayit)
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Ürün ve hat bazında fire oranını ve verim kaybı maliyetini hesaplar.")
    p.add_argument("--kayitlar", type=Path, default=ORNEK / "uretim_kayitlari.csv",
                   help="Tarih, Hat, Ürün, Hammadde Girişi (kg), Mamul Çıkışı (kg), Geri Kazanılan (kg), Birim Maliyet (TL/kg), Standart Fire %%")
    p.add_argument("--standart", help="Dosyada standart fire yoksa kullanılacak oran %% (varsayılan: yok)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "fire_verim_analizi.xlsx")
    a = p.parse_args(argv)
    if not a.kayitlar.exists():
        print(f"[X] Dosya bulunamadı: {a.kayitlar}")
        return 1
    std = para(a.standart) if a.standart else None
    if a.standart and (std is None or not 0 <= std < 100):
        print("[X] --standart 0 ile 100 arasında bir yüzde olmalı")
        return 2
    try:
        r = calistir(a.kayitlar, a.cikti, std)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {t['kayit']} kayıt · fire %{t['fire_oran']} · verim %{t['verim']}" + (f" · fire maliyeti {t['fire_tl']:,.2f}" if t["fire_tl"] is not None else "")
          + (f" · standart üstü {t['ustu_tl']:,.2f}" if t["ustu_tl"] is not None else ""))
    for x in r["urun_hat"][:3]:
        if x["ustu_tl"]:
            print(f"[!] {x['ad']}: fire %{x['fire_oran']} · verim kaybı maliyeti {x['ustu_tl']:,.2f}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
