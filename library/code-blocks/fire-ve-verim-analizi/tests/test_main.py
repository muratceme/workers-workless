import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def k(giris=1000, cikis=940, geri=20, maliyet=50, standart=3, urun="A", hat="H1"):
    return {"tarih": None, "hat": hat, "urun": urun, "giris": D(giris), "cikis": D(cikis), "geri": D(geri), "maliyet": None if maliyet is None else D(maliyet),
            "standart": None if standart is None else D(standart)}


class HesapTesti(unittest.TestCase):
    def test_kayit(self):
        x = main.hesapla(k())
        self.assertEqual((x["fire"], x["fire_oran"], x["verim"], x["fire_tl"], x["ustu"], x["ustu_tl"]), (D(40), D("4.00"), D("94.00"), D("2000.00"), D(10), D("500.00")))
        x = main.hesapla(k(1200, 1150, 14))                                  # fire 36 = tam standart
        self.assertEqual((x["fire_oran"], x["ustu"], x["ustu_tl"]), (D("3.00"), D(0), D("0.00")))
        x = main.hesapla(k(maliyet=None, standart=None))
        self.assertEqual((x["fire"], x["fire_tl"], x["ustu"], x["ustu_tl"]), (D(40), None, None, None))

    def test_gruplar(self):
        r = main.analiz([k(), k(1200, 1150, 14), k(2000, 1800, 50, 40, 5, "B", "H2"), k(800, 700, 20, 50, 3, "A", "H2")])
        t = r["toplam"]
        self.assertEqual((t["giris"], t["fire"], t["fire_oran"], t["verim"], t["fire_tl"], t["ustu"], t["ustu_tl"], t["asan"]),
                         (D(5000), D(306), D("6.12"), D("91.80"), D("13800.00"), D(116), D("5300.00"), 3))
        self.assertEqual([(x["ad"], x["fire"], x["fire_oran"], x["ustu_tl"]) for x in r["urun"]], [("A", D(156), D("5.20"), D("3300.00")), ("B", D(150), D("7.50"), D("2000.00"))])
        self.assertEqual([(x["ad"], x["ustu_tl"]) for x in r["hat"]], [("H2", D("4800.00")), ("H1", D("500.00"))])
        self.assertEqual(r["urun_hat"][0]["ad"], "A · H2")
        with self.assertRaises(ValueError):
            main.analiz([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "f.xlsx"
            r = main.calistir(ORNEK / "uretim_kayitlari.csv", cikti)
            t = r["toplam"]
            self.assertEqual((t["kayit"], t["giris"], t["cikis"], t["geri"], t["fire"], t["fire_oran"], t["verim"], t["fire_tl"], t["ustu"], t["ustu_tl"]),
                             (5, D(7000), D(6490), D(124), D(386), D("5.51"), D("92.71"), D("17000.00"), D(116), D("5300.00")))
            self.assertEqual({x["ad"]: (x["fire_oran"], x["fire_tl"]) for x in r["urun"]}, {"Kapak A": (D("5.20"), D("7800.00")), "Gövde B": (D("5.75"), D("9200.00"))})
            self.assertEqual({x["ad"]: (x["fire_oran"], x["ustu_tl"]) for x in r["hat"]}, {"Hat 1": (D("3.45"), D("500.00")), "Hat 2": (D("6.46"), D("4800.00"))})
            self.assertEqual([u["tur"] for u in r["uyarilar"]], ["Çıkış girişten büyük"])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Ürün Bazında", "Hat Bazında", "Ürün-Hat", "Kayıtlar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "k.csv"
            y.write_text("Ürün;Giriş (kg);Çıkış (kg);Birim Maliyet;Standart Fire %\nA;1.000;900;pahalı;150\nB;0;5;;\nC;yok;5;;\nD;100;90;10;\n", encoding="utf-8")
            l, uy = main.oku(y, D(4))
            self.assertEqual([(x["urun"], x["giris"], x["maliyet"], x["standart"], x["hat"]) for x in l], [("A", D(1000), None, D(4), "Belirtilmemiş"), ("D", D(100), D(10), D(4), "Belirtilmemiş")])
            self.assertEqual(sorted(u["tur"] for u in uy), ["Birim maliyet okunamadı", "Miktar okunamadı", "Miktar okunamadı", "Standart fire okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            self.assertEqual(main.main(["--kayitlar", str(y), "--cikti", c, "--standart", "4"]), 0)
            self.assertEqual(main.main(["--kayitlar", str(y), "--cikti", c, "--standart", "çok"]), 2)
            self.assertEqual(main.main(["--kayitlar", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Ürün;Giriş;Çıkış\n", encoding="utf-8")
            self.assertEqual(main.main(["--kayitlar", str(y), "--cikti", c]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
