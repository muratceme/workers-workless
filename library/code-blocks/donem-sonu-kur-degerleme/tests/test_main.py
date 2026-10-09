import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
T = date(2026, 12, 31)


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "k.xlsx"
        cls.s = main.calistir(ORNEK / "bakiyeler.csv", ORNEK / "kurlar.csv", cls.cikti, T)
        cls.h = {h.kod: h for h in cls.s["hesaplar"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_kur_farklari(self):
        self.assertEqual(self.h["102.01.01"].fark, D("41012.50"))          # 48.250 × 41,45 − 1.958.950
        self.assertEqual(self.h["300.01.0001"].fark, D("-237000.00"))      # pasif: borç arttı → zarar
        self.assertEqual(self.h["120.01.0060"].fark, D("-1250.40"))        # döviz 0, TL kalıntısı
        self.assertEqual((self.s["kar"], self.s["zarar"]), (D("126279.50"), D("279576.40")))

    def test_hariç_ve_uyarilar(self):
        self.assertEqual(self.h["159.01.0003"].durum, "Değerlemeye alınmadı")
        self.assertEqual(self.h["340.01.0007"].durum, "Değerlemeye alınmadı")
        self.assertEqual(self.h["120.01.0072"].durum, "Kur yok")
        u = {(x["tur"], x["hesap"]) for x in self.s["uyarilar"]}
        self.assertTrue({("Kur yok", "120.01.0072"), ("TL kalıntısı", "120.01.0060"), ("Ters bakiye", "136.01.0001"), ("Avans hesabı", "159.01.0003")} <= u)

    def test_yevmiye_ve_excel(self):
        yv = self.s["yevmiye"]
        self.assertEqual(sum(x["borc"] for x in yv), sum(x["alacak"] for x in yv))
        self.assertEqual([x["hesap"] for x in yv[-2:]], ["646", "656"])
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Değerleme", "Yevmiye Önerisi", "Döviz Özeti", "Uyarılar"])
        self.assertEqual(wb["Yevmiye Önerisi"]["F14"].value, "Dengede")


class KuralTesti(unittest.TestCase):
    def test_avans_dahil_ozel_kur_sapma(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t / "b.csv").write_text("Hesap Kodu;Döviz;Döviz Bakiye;Borç/Alacak;TL Bakiye;Kur\n159.01;USD;100;B;4000;\n102.05;USD;100;B;3000;\n"
                                     "320.09;EUR;-100;;-4700;48,50\n400.01;USD;100;B;4100;\n", encoding="utf-8")
            s = main.calistir(t / "b.csv", ORNEK / "kurlar.csv", t / "o.xlsx", T, avans_dahil=True, kar_hesabi="646.01", zarar_hesabi="656.01")
            h = {x.kod: x for x in s["hesaplar"]}
            self.assertEqual(h["159.01"].fark, D("145.00"))
            self.assertEqual(h["320.09"].fark, D("-150.00"))               # satırdaki özel kur 48,50
            turler = {(u["tur"], u["hesap"]) for u in s["uyarilar"]}
            self.assertIn(("Kayıtlı kur sapması", "102.05"), turler)         # 30 / 41,45 → %28
            self.assertIn(("Ters bakiye", "400.01"), turler)
            self.assertIn("656.01", [x["hesap"] for x in s["yevmiye"]])

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--tarih", "2026"]), 2)
            self.assertEqual(main.main(["--kurlar", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
