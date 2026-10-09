"""
Dönem Sonu Kur Değerleme — Workers / Workless kod bloğu
Muhasebe › Genel Muhasebe Uzmanı

Döviz cinsinden hesap bakiyelerini dönem sonu kurlarıyla değerler ve kur farkı kayıtlarını hesaplar:
  - Değerlenmiş TL = döviz bakiye × değerleme kuru (hesap satırında "Kur" verilirse o kullanılır).
  - Kur farkı = değerlenmiş TL − kayıtlı TL (borç bakiyesi +, alacak bakiyesi −).
    Fark > 0: hesap borç / 646 Kambiyo Kârları alacak. Fark < 0: 656 Kambiyo Zararları borç / hesap alacak.
  - Döviz bakiyesi sıfır olduğu hâlde TL bakiyesi kalan hesapta TL bakiyesi kur farkıyla kapatılır.
  - Sipariş avansı hesapları (159, 179, 340, 349, 440) parasal kalem sayılmadığından varsayılan olarak
    değerlemeye alınmaz (--avanslari-dahil-et).
  - Kontroller: kuru verilmemiş döviz, beklenenin tersine bakiye veren hesap (ör. alacak bakiyeli 120),
    kayıtlı kurun değerleme kurundan %20'den fazla sapması (veri hatası olabilir).
Rapor: değerleme tablosu, yevmiye kaydı önerisi, döviz özeti, uyarılar. İnternete bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: 13 hesap, 31.12.2026
    python main.py --bakiyeler bakiyeler.xlsx --kurlar kurlar.csv --tarih 31.12.2026
    python main.py --bakiyeler mizan_doviz.xlsx --kurlar kurlar.csv --kar-hesabi 646 --zarar-hesabi 656
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
K2 = Decimal("0.01")
AVANS_HESAPLARI = ("159", "179", "340", "349", "440")
# Tek Düzen Hesap Planı'na göre normalde borç (aktif) / alacak (pasif) bakiye veren ana hesap grupları
AKTIF_ON_EK = ("10", "11", "12", "13", "22", "23", "24", "25")
PASIF_ON_EK = ("30", "32", "33", "34", "40", "42", "43", "44")

BAKIYE_SUTUNLARI = {"kod": ("hesap kodu", "hesap no", "kod"), "ad": ("hesap adi", "aciklama", "unvan"), "doviz": ("doviz", "doviz cinsi", "para birimi"),
                    "dbakiye": ("doviz bakiye", "doviz bakiyesi", "doviz tutari"), "ba": ("borc alacak", "b a", "bakiye yonu"),
                    "tl": ("tl bakiye", "tl bakiyesi", "kayitli tl", "tl tutari"), "kur": ("kur", "degerleme kuru")}
KUR_SUTUNLARI = {"doviz": ("doviz", "doviz cinsi", "para birimi", "kod"), "kur": ("kur", "degerleme kuru", "doviz alis", "doviz alis kuru"),
                 "kaynak": ("kaynak", "aciklama")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tl(x: Decimal | None) -> str:
    return "—" if x is None else f"{x.quantize(K2, ROUND_HALF_UP):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


# ----------------------------------------------------------------------------
# Veri ve hesap
# ----------------------------------------------------------------------------

@dataclass
class Hesap:
    kod: str
    ad: str
    doviz: str
    dbakiye: Decimal          # işaretli: borç +, alacak −
    tl: Decimal               # işaretli
    kur_ozel: Decimal | None = None
    kur: Decimal | None = None
    degerlenmis: Decimal | None = None
    durum: str = ""
    notlar: list = field(default_factory=list)

    @property
    def fark(self) -> Decimal | None:
        return None if self.degerlenmis is None else self.degerlenmis - self.tl

    @property
    def kayitli_kur(self) -> Decimal | None:
        return abs(self.tl / self.dbakiye) if self.dbakiye else None


def oku_bakiyeler(yol: Path) -> tuple[list[Hesap], list[dict]]:
    sonuc, uy = [], []
    for r in kayitlar(yol, BAKIYE_SUTUNLARI, ("kod", "doviz", "dbakiye", "tl")):
        kod = metin(r.get("kod"))
        if not kod:
            continue
        d, t = para(r.get("dbakiye")), para(r.get("tl"))
        if d is None or t is None:
            uy.append({"onem": "Orta", "tur": "Okunamayan satır", "hesap": kod, "aciklama": f"Satır {r['_satir']}: döviz veya TL bakiyesi okunamadı"})
            continue
        ba = katla(r.get("ba"))
        if ba.startswith("a"):
            d, t = -abs(d), -abs(t)
        elif ba.startswith("b"):
            d, t = abs(d), abs(t)
        sonuc.append(Hesap(kod, metin(r.get("ad")), metin(r.get("doviz")).upper(), d, t, para(r.get("kur"))))
    return sonuc, uy


def oku_kurlar(yol: Path) -> dict[str, tuple[Decimal, str]]:
    k = {}
    for r in kayitlar(yol, KUR_SUTUNLARI, ("doviz", "kur")):
        v = para(r.get("kur"))
        if metin(r.get("doviz")) and v:
            k[metin(r["doviz"]).upper()] = (v, metin(r.get("kaynak")))
    return k


def degerle(hesaplar: list[Hesap], kurlar: dict[str, tuple[Decimal, str]], avans_dahil: bool = False, sapma: Decimal = Decimal("0.20")) -> list[dict]:
    uyarilar = []

    def uyar(onem, tur, h, aciklama):
        uyarilar.append({"onem": onem, "tur": tur, "hesap": h.kod if h else "", "aciklama": aciklama})

    for h in hesaplar:
        if h.doviz in ("", "TRY", "TL"):
            h.durum = "Değerlemeye alınmadı"
            h.notlar.append("TL hesap")
            continue
        if h.kod.startswith(AVANS_HESAPLARI) and not avans_dahil:
            h.durum = "Değerlemeye alınmadı"
            h.notlar.append("Sipariş avansı: parasal kalem değil (--avanslari-dahil-et)")
            uyar("Bilgi", "Avans hesabı", h, f"{h.ad}: değerlemeye alınmadı. Avansın niteliğine göre (iadesi söz konusu ise) değerlendirin")
            continue
        h.kur = h.kur_ozel or (kurlar.get(h.doviz) or (None,))[0]
        if h.kur is None:
            h.durum = "Kur yok"
            uyar("Yüksek", "Kur yok", h, f"{h.doviz} için değerleme kuru verilmedi; hesap değerlenmedi")
            continue
        h.degerlenmis = (h.dbakiye * h.kur).quantize(K2, ROUND_HALF_UP)
        h.durum = "Kur farkı yok" if h.fark == 0 else "Kâr" if h.fark > 0 else "Zarar"
        if h.dbakiye == 0 and h.tl != 0:
            h.notlar.append("Döviz bakiyesi sıfır; TL kalıntısı kur farkıyla kapatılır")
            uyar("Orta", "TL kalıntısı", h, f"Döviz bakiyesi 0, TL bakiyesi {tl(h.tl)}; kapanış kaydı önerildi. Kalıntının nedenini kontrol edin")
        ana = h.kod[:2]
        if (ana in AKTIF_ON_EK and h.dbakiye < 0) or (ana in PASIF_ON_EK and h.dbakiye > 0):
            h.notlar.append("Ters bakiye")
            uyar("Orta", "Ters bakiye", h, f"{h.ad}: {'alacak' if h.dbakiye < 0 else 'borç'} bakiyesi veriyor; fazla tahsilat / ödeme olabilir, "
                 "sınıflandırmayı kontrol edin")
        if (h.dbakiye > 0) != (h.tl > 0) and h.dbakiye and h.tl:
            uyar("Yüksek", "Yön uyuşmazlığı", h, f"Döviz bakiyesi ile TL bakiyesinin yönü farklı ({tl(h.dbakiye)} {h.doviz} / {tl(h.tl)} TL)")
        kk = h.kayitli_kur
        if kk and abs(kk - h.kur) / h.kur > sapma:
            uyar("Orta", "Kayıtlı kur sapması", h, f"Kayıtlı ortalama kur {kk:.4f}, değerleme kuru {h.kur:.4f} (%{abs(kk - h.kur) / h.kur * 100:.0f} fark); "
                 "döviz veya TL bakiyesini kontrol edin")
    return uyarilar


def yevmiye(hesaplar: list[Hesap], kar_hesabi: str, zarar_hesabi: str) -> list[dict]:
    satirlar = []
    for h in hesaplar:
        f = h.fark
        if not f:
            continue
        if f > 0:
            satirlar.append({"hesap": h.kod, "ad": h.ad, "borc": f, "alacak": Decimal(0), "aciklama": f"{h.doviz} kur değerlemesi ({tl(h.kur)})"})
        else:
            satirlar.append({"hesap": h.kod, "ad": h.ad, "borc": Decimal(0), "alacak": -f, "aciklama": f"{h.doviz} kur değerlemesi ({tl(h.kur)})"})
    kar = sum((h.fark for h in hesaplar if h.fark and h.fark > 0), Decimal(0))
    zarar = sum((-h.fark for h in hesaplar if h.fark and h.fark < 0), Decimal(0))
    if kar:
        satirlar.append({"hesap": kar_hesabi, "ad": "Kambiyo Kârları", "borc": Decimal(0), "alacak": kar, "aciklama": "Dönem sonu kur değerlemesi"})
    if zarar:
        satirlar.append({"hesap": zarar_hesabi, "ad": "Kambiyo Zararları", "borc": zarar, "alacak": Decimal(0), "aciklama": "Dönem sonu kur değerlemesi"})
    return satirlar


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE", "Kâr": "E3F4E1", "Zarar": "FDE2E1", "Kur yok": "FDE2E1",
        "Değerlemeye alınmadı": "EEEEEE", "Kur farkı yok": "FFFFFF"}
KONTROL = PatternFill("solid", fgColor="FFF4CE")
UST = Alignment(vertical="top", wrap_text=True)
PF = "#,##0.00"


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def _f(x):
    return None if x is None else float(x)


def rapor_yaz(cikti: Path, s: dict) -> None:
    wb = Workbook()
    dg = wb.active
    dg.title = "Değerleme"
    _baslik(dg, ["Hesap Kodu", "Hesap Adı", "Döviz", "Döviz Bakiye", "B/A", "Kayıtlı TL", "Kayıtlı Ort. Kur", "Değerleme Kuru", "Değerlenmiş TL",
                 "Kur Farkı", "Durum", "Not", "Kontrol"], (14, 32, 6, 14, 5, 15, 10, 10, 15, 13, 13, 40, 14))
    for h in s["hesaplar"]:
        dg.append([h.kod, h.ad, h.doviz, float(abs(h.dbakiye)), "A" if h.dbakiye < 0 or (h.dbakiye == 0 and h.tl < 0) else "B", float(abs(h.tl)),
                   _f(h.kayitli_kur), _f(h.kur), _f(abs(h.degerlenmis)) if h.degerlenmis is not None else None, _f(h.fark), h.durum, "; ".join(h.notlar), ""])
        r = dg.max_row
        for j in (4, 6, 9, 10):
            dg.cell(r, j).number_format = PF
        for j in (7, 8):
            dg.cell(r, j).number_format = "0.0000"
        dg.cell(r, 11).fill = PatternFill("solid", fgColor=RENK[h.durum])
        dg.cell(r, 13).fill = KONTROL
    dg.append([])
    dg.append(["Toplam kur farkı (net)", "", "", None, "", None, None, None, None, float(s["net"])])
    dg.cell(dg.max_row, 10).number_format = PF
    dg.cell(dg.max_row, 10).font = Font(bold=True)
    dg.auto_filter.ref = f"A1:M{len(s['hesaplar']) + 1}"

    yv = wb.create_sheet("Yevmiye Önerisi")
    _baslik(yv, ["Tarih", "Hesap Kodu", "Hesap Adı", "Borç", "Alacak", "Açıklama"], (11, 14, 32, 15, 15, 40))
    for x in s["yevmiye"]:
        yv.append([s["tarih"], x["hesap"], x["ad"], float(x["borc"]) or None, float(x["alacak"]) or None, x["aciklama"]])
        yv.cell(yv.max_row, 1).number_format = "DD.MM.YYYY"
        for j in (4, 5):
            yv.cell(yv.max_row, j).number_format = PF
    tb = sum((x["borc"] for x in s["yevmiye"]), Decimal(0))
    ta = sum((x["alacak"] for x in s["yevmiye"]), Decimal(0))
    yv.append(["", "", "Toplam", float(tb), float(ta), "Dengede" if tb == ta else "DENGESİZ"])
    for j in (4, 5):
        yv.cell(yv.max_row, j).number_format = PF
        yv.cell(yv.max_row, j).font = Font(bold=True)
    yv.append([])
    yv.append(["Not: Kayıt önerisidir; muhasebe programına aktarmadan önce kontrol edin. Yatırım dönemine ait kredi kur farklarının ilgili "
               "varlığın maliyetine eklenmesi gibi özel durumlar ayrıca değerlendirilmelidir."])

    oz = wb.create_sheet("Döviz Özeti")
    _baslik(oz, ["Döviz", "Değerleme Kuru", "Kur Kaynağı", "Net Döviz Pozisyonu", "Kayıtlı TL (net)", "Değerlenmiş TL (net)", "Kur Farkı"],
            (7, 12, 50, 16, 16, 16, 14))
    g = defaultdict(lambda: [Decimal(0), Decimal(0), Decimal(0)])
    for h in s["hesaplar"]:
        if h.degerlenmis is not None:
            g[h.doviz][0] += h.dbakiye
            g[h.doviz][1] += h.tl
            g[h.doviz][2] += h.degerlenmis
    for d, (pos, kayit, deg) in sorted(g.items()):
        k, kaynak = s["kurlar"].get(d, (None, "hesap satırındaki kur"))
        oz.append([d, _f(k), kaynak, float(pos), float(kayit), float(deg), float(deg - kayit)])
        oz.cell(oz.max_row, 2).number_format = "0.0000"
        for j in (4, 5, 6, 7):
            oz.cell(oz.max_row, j).number_format = PF
    oz.append([])
    oz.append(["Net pozisyon: borç bakiyeli (varlık) +, alacak bakiyeli (borç) −. Pozitif pozisyonda kur artışı kâr, negatifte zarar yazar."])

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Hesap", "Açıklama"], (9, 20, 14, 90))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["hesap"], u["aciklama"]])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(bakiye_yolu: Path, kur_yolu: Path, cikti: Path, tarih_: date, *, avans_dahil: bool = False, kar_hesabi: str = "646",
             zarar_hesabi: str = "656") -> dict:
    hesaplar, uyarilar = oku_bakiyeler(bakiye_yolu)
    if not hesaplar:
        raise ValueError(f"{bakiye_yolu.name}: döviz hesabı bulunamadı")
    kurlar = oku_kurlar(kur_yolu)
    uyarilar += degerle(hesaplar, kurlar, avans_dahil)
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    uyarilar.sort(key=lambda u: (sira[u["onem"]], u["tur"], u["hesap"]))
    yv = yevmiye(hesaplar, kar_hesabi, zarar_hesabi)
    net = sum((h.fark for h in hesaplar if h.fark), Decimal(0))
    s = {"hesaplar": hesaplar, "kurlar": kurlar, "uyarilar": uyarilar, "yevmiye": yv, "net": net, "tarih": tarih_,
         "kar": sum((h.fark for h in hesaplar if h.fark and h.fark > 0), Decimal(0)), "zarar": sum((-h.fark for h in hesaplar if h.fark and h.fark < 0), Decimal(0))}
    rapor_yaz(cikti, s)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Döviz cinsinden hesapları dönem sonu kurlarıyla değerler, kur farkı kaydı önerir.")
    p.add_argument("--bakiyeler", type=Path, default=ORNEK / "bakiyeler.csv", help="Hesap Kodu, Hesap Adı, Döviz, Döviz Bakiye, Borç/Alacak, TL Bakiye, [Kur]")
    p.add_argument("--kurlar", type=Path, default=ORNEK / "kurlar.csv", help="Döviz, Kur, Kaynak")
    p.add_argument("--tarih", help="Değerleme tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 31.12.2026)")
    p.add_argument("--avanslari-dahil-et", action="store_true", help="159, 179, 340, 349, 440 avans hesaplarını da değerle")
    p.add_argument("--kar-hesabi", default="646", help="Kur farkı kârı hesabı (varsayılan 646)")
    p.add_argument("--zarar-hesabi", default="656", help="Kur farkı zararı hesabı (varsayılan 656)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kur_degerleme.xlsx")
    a = p.parse_args(argv)
    for y in (a.bakiyeler, a.kurlar):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    t = tarih(a.tarih) if a.tarih else (date(2026, 12, 31) if a.bakiyeler == ORNEK / "bakiyeler.csv" else date.today())
    if not t:
        print("[X] --tarih GG.AA.YYYY biçiminde olmalı")
        return 2
    try:
        s = calistir(a.bakiyeler, a.kurlar, a.cikti, t, avans_dahil=a.avanslari_dahil_et, kar_hesabi=a.kar_hesabi, zarar_hesabi=a.zarar_hesabi)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(s['hesaplar'])} hesap · kur farkı kârı {tl(s['kar'])} TL, zararı {tl(s['zarar'])} TL, net {tl(s['net'])} TL")
    for u in s["uyarilar"]:
        if u["onem"] != "Bilgi":
            print(f"[!] {u['hesap']} {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
