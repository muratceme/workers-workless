"""
Temsilci İhtiyacı Hesaplama (Erlang C) — Workers / Workless kod bloğu
Müşteri Hizmetleri ve Çağrı Merkezi › Çağrı Merkezi Yöneticisi

Aralık (30/60 dk) bazında çağrı tahmininden, ortalama görüşme süresi (AHT), hedef servis seviyesi
(ör. çağrıların %80'i 20 saniyede) ve en yüksek doluluk oranına göre gereken temsilci sayısını
Erlang C modeliyle hesaplar; izin, eğitim, mola gibi kayıplar (shrinkage) için planlanacak kişi
sayısını da verir. İnternete bağlanmaz.

Kullanım:
    python main.py                                              # örnek günlük tahminle dener
    python main.py --girdi cagri_tahmini.xlsx --aht 180 --sl 80 --hedef-sn 20 --kayip 30
    python main.py --cagri 100 --aralik 30 --aht 180            # tek aralık
"""
from __future__ import annotations

import argparse
import csv
import math
import sys
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent


def erlang_c(trafik: float, n: int) -> float:
    """Bekleme olasılığı P(W>0). Erlang B özyinelemesiyle (büyük n için de kararlı)."""
    if n <= trafik:
        return 1.0
    b = 1.0
    for k in range(1, n + 1):
        b = trafik * b / (k + trafik * b)
    return n * b / (n - trafik * (1 - b))


@dataclass
class Sonuc:
    cagri: float
    aht: float
    trafik: float
    temsilci: int
    bekleme_olasiligi: float
    servis_seviyesi: float
    ort_cevap_suresi: float
    doluluk: float
    planlanan: int


def hesapla(cagri: float, aralik_dk: float, aht_sn: float, sl_hedef: float = 0.80, hedef_sn: float = 20,
            maks_doluluk: float = 0.85, kayip: float = 0.0, n: int | None = None) -> Sonuc:
    """n verilirse o temsilci sayısının performansını, verilmezse hedefleri sağlayan en küçük n'yi döndürür."""
    trafik = cagri * aht_sn / (aralik_dk * 60)          # Erlang
    if cagri <= 0:
        return Sonuc(cagri, aht_sn, 0.0, 0, 0.0, 1.0, 0.0, 0.0, 0)

    def performans(m: int):
        pw = erlang_c(trafik, m)
        sl = 1 - pw * math.exp(-(m - trafik) * hedef_sn / aht_sn) if m > trafik else 0.0
        asa = pw * aht_sn / (m - trafik) if m > trafik else math.inf
        return pw, sl, asa, trafik / m

    if n is None:
        n = max(1, math.floor(trafik) + 1)
        while True:
            pw, sl, asa, dol = performans(n)
            if sl >= sl_hedef and dol <= maks_doluluk:
                break
            n += 1
    pw, sl, asa, dol = performans(n)
    planlanan = math.ceil(n / (1 - kayip)) if kayip < 1 else n
    return Sonuc(cagri, aht_sn, trafik, n, pw, sl, asa, dol, planlanan)


# ----------------------------------------------------------------------------
# Girdi / çıktı
# ----------------------------------------------------------------------------

def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def sayi(x) -> float | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return float(x)
    return float(str(x).strip().replace(",", "."))


def tahmin_oku(yol: Path) -> list[dict]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ayirici = csv.Sniffer().sniff(metin.splitlines()[0], delimiters=",;\t").delimiter
        satirlar = list(csv.reader(metin.splitlines(), delimiter=ayirici))
    satirlar = [r for r in satirlar if any(c not in (None, "") for c in r)]
    b = [kucuk(x) for x in satirlar[0]]
    bul = lambda *adlar: next((i for i, x in enumerate(b) if x in adlar), None)  # noqa: E731
    i_gun, i_ara = bul("gün", "tarih"), bul("aralık", "saat", "aralık başlangıcı", "zaman")
    i_cagri, i_aht = bul("çağrı", "çağrı sayısı", "gelen çağrı", "tahmini çağrı"), bul("aht", "aht (sn)", "ortalama görüşme süresi (sn)")
    if i_cagri is None or i_ara is None:
        raise SystemExit(f"Gerekli sütunlar: Aralık (veya Saat) ve Çağrı Sayısı. Başlıklar: {satirlar[0]}")
    return [{"gun": r[i_gun] if i_gun is not None else "", "aralik": r[i_ara], "cagri": sayi(r[i_cagri]) or 0,
             "aht": sayi(r[i_aht]) if i_aht is not None else None} for r in satirlar[1:]]


BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")


def rapor_yaz(satirlar: list[tuple[dict, Sonuc]], cikti: Path, ayarlar: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Aralık Planı"
    ws.append(["Gün", "Aralık", "Çağrı", "AHT (sn)", "Trafik (Erlang)", "Gereken Temsilci", "Planlanan (kayıp dahil)",
               "Servis Seviyesi", "Bekleme Olasılığı", "Ort. Cevap Süresi (sn)", "Doluluk"])
    for h in ws[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for t, s in satirlar:
        ws.append([t["gun"], t["aralik"], s.cagri, s.aht, round(s.trafik, 2), s.temsilci, s.planlanan,
                   s.servis_seviyesi, s.bekleme_olasiligi, round(s.ort_cevap_suresi, 1) if math.isfinite(s.ort_cevap_suresi) else None,
                   s.doluluk])
        for c in (8, 9, 11):
            ws.cell(ws.max_row, c).number_format = "0.0%"
    for j, g in enumerate((12, 10, 9, 9, 15, 17, 22, 15, 17, 21, 10), 1):
        ws.column_dimensions[get_column_letter(j)].width = g
    ws.freeze_panes = "C2"
    son = ws.max_row
    if son > 1:
        g = BarChart()
        g.title, g.height, g.width = "Aralık bazında gereken temsilci", 9, 24
        g.add_data(Reference(ws, min_col=6, max_col=7, min_row=1, max_row=son), titles_from_data=True)
        g.set_categories(Reference(ws, min_col=2, min_row=2, max_row=son))
        c = LineChart()
        c.add_data(Reference(ws, min_col=3, min_row=1, max_row=son), titles_from_data=True)
        c.y_axis.axId, c.y_axis.title = 200, "Çağrı"
        g.y_axis.title = "Temsilci"
        c.y_axis.crosses = "max"
        g += c                    # openpyxl: grafik birleştirme yalnızca += ile
        ws.add_chart(g, "M2")

    oz = wb.create_sheet("Özet")
    toplam_cagri = sum(s.cagri for _, s in satirlar)
    saat = ayarlar["aralik"] / 60
    for r in [
        ["Hedef servis seviyesi", f"%{ayarlar['sl'] * 100:g} çağrı {ayarlar['hedef_sn']:g} saniyede"],
        ["En yüksek doluluk", f"%{ayarlar['maks_doluluk'] * 100:g}"], ["Kayıp (shrinkage)", f"%{ayarlar['kayip'] * 100:g}"],
        ["Aralık uzunluğu (dk)", ayarlar["aralik"]], ["Toplam çağrı", toplam_cagri],
        ["En yoğun aralıkta gereken temsilci", max((s.temsilci for _, s in satirlar), default=0)],
        ["En yoğun aralıkta planlanacak kişi", max((s.planlanan for _, s in satirlar), default=0)],
        ["Toplam temsilci-saat (gereken)", round(sum(s.temsilci for _, s in satirlar) * saat, 1)],
        ["Toplam temsilci-saat (planlanan)", round(sum(s.planlanan for _, s in satirlar) * saat, 1)],
        ["Model", "Erlang C: sonsuz sabır (çağrı terk etmez), Poisson geliş, üstel görüşme süresi varsayımları"],
    ]:
        oz.append(r)
    oz.column_dimensions["A"].width = 36
    oz.column_dimensions["B"].width = 90
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description="Erlang C ile aralık bazında temsilci ihtiyacı.")
    ap.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "cagri_tahmini.csv")
    ap.add_argument("--cagri", type=float, help="Tek aralık için çağrı sayısı")
    ap.add_argument("--aralik", type=float, default=30, help="Aralık uzunluğu (dk): 15, 30 veya 60")
    ap.add_argument("--aht", type=float, default=180, help="Ortalama görüşme süresi (sn); girdide AHT sütunu varsa o kullanılır")
    ap.add_argument("--sl", type=float, default=80, help="Hedef servis seviyesi (%%)")
    ap.add_argument("--hedef-sn", type=float, default=20, help="Hedef cevaplama süresi (sn)")
    ap.add_argument("--maks-doluluk", type=float, default=85, help="En yüksek temsilci doluluğu (%%)")
    ap.add_argument("--kayip", type=float, default=30, help="Kayıp oranı / shrinkage (%%): izin, eğitim, mola, devamsızlık")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "temsilci_ihtiyaci.xlsx")
    a = ap.parse_args(argv)
    ayar = {"aralik": a.aralik, "sl": a.sl / 100, "hedef_sn": a.hedef_sn, "maks_doluluk": a.maks_doluluk / 100, "kayip": a.kayip / 100}

    if a.cagri is not None:
        s = hesapla(a.cagri, a.aralik, a.aht, ayar["sl"], a.hedef_sn, ayar["maks_doluluk"], ayar["kayip"])
        print((f"[OK] {a.cagri:g} çağrı / {a.aralik:g} dk, AHT {a.aht:g} sn → trafik {s.trafik:.2f} Erlang · "
              f"{s.temsilci} temsilci (SL %{s.servis_seviyesi * 100:.1f}, ASA {s.ort_cevap_suresi:.1f} sn, doluluk %{s.doluluk * 100:.1f}) · "
              f"kayıp dahil {s.planlanan} kişi").replace(".", ","))
        return
    tahmin = tahmin_oku(a.girdi)
    satirlar = [(t, hesapla(t["cagri"], a.aralik, t["aht"] or a.aht, ayar["sl"], a.hedef_sn, ayar["maks_doluluk"], ayar["kayip"]))
                for t in tahmin]
    rapor_yaz(satirlar, a.cikti, ayar)
    tepe = max(satirlar, key=lambda x: x[1].temsilci)
    print(f"[OK] {len(satirlar)} aralık · en yoğun {tepe[0]['aralik']}: {tepe[1].temsilci} temsilci (kayıp dahil {tepe[1].planlanan})")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
