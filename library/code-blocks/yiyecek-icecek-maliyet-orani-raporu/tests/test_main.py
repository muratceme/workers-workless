import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def s(dep, tur, acilis, alim, kapanis, satis, hedef=None, tr_g=0, tr_c=0, personel=0, ikram=0, onceki=None):
    return {"dep": dep, "tur": tur, "acilis": D(acilis), "alim": D(alim), "tr_giris": D(tr_g), "tr_cikis": D(tr_c), "kapanis": D(kapanis), "personel": D(personel), "ikram": D(ikram),
            "satis": D(satis), "hedef": None if hedef is None else D(hedef), "onceki": None if onceki is None else D(onceki)}


class AnalizTesti(unittest.TestCase):
    def test_oranlar(self):
        l = [s("Restoran", "Yiyecek", 100, 500, 120, 1500, 30, tr_c=30, personel=40, ikram=10, onceki=28), s("Bar", "İçecek", 50, 200, 40, 1000, 20, tr_g=30, onceki=22),
             s("Banket", "Yiyecek", 20, 100, 30, 600, 25, onceki=20)]
        r = main.analiz(l, gun=30)
        x = {a["dep"]: a for a in r["satirlar"]}
        # Restoran: 100 + 500 − 30 − 120 = 450 tüketim; 450 − 40 − 10 = 400; 400 ÷ 1500 = %26,67
        self.assertEqual((x["Restoran"]["tuketim"], x["Restoran"]["maliyet"], x["Restoran"]["oran"], x["Restoran"]["fark"], x["Restoran"]["degisim"], x["Restoran"]["stok_gun"]),
                         (D(450), D(400), D("26.67"), D("-3.33"), D("-1.33"), D("7.3")))
        self.assertEqual((x["Bar"]["maliyet"], x["Bar"]["oran"], x["Bar"]["fark"]), (D(240), D("24.00"), D("4.00")))
        self.assertEqual((x["Banket"]["oran"], x["Banket"]["fark"]), (D("15.00"), D("-10.00")))
        y = r["turler"][0]
        # Yiyecek toplamı: maliyet 490 ÷ satış 2100; hedef (30×1500 + 25×600) ÷ 2100 = 28,57
        self.assertEqual((y["tur"], y["maliyet"], y["oran"], y["hedef"], y["onceki"]), ("Yiyecek", D(490), D("23.33"), D("28.57"), D("25.71")))
        self.assertEqual((r["genel"]["maliyet"], r["genel"]["oran"]), (D(730), D("23.55")))
        self.assertEqual([a["dep"] for a in r["satirlar"]], ["Banket", "Restoran", "Bar"])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Bilgi", "Maliyet oranı hedefin çok altında", "Banket · Yiyecek"), ("Orta", "Maliyet oranı hedefin üstünde", "Bar · İçecek")])
        self.assertEqual(main.analiz(l, D(5))["uyarilar"][0]["kim"], "Banket · Yiyecek")

    def test_ozel_durumlar(self):
        r = main.analiz([s("A", "Yiyecek", 10, 0, 20, 0, 30), s("B", "İçecek", 10, 10, -5, 100, tr_g=5)])
        self.assertEqual(sorted(u["tur"] for u in r["uyarilar"]), ["Negatif stok / maliyet", "Negatif stok / maliyet", "Satış yok", "Transferler denk değil"])
        self.assertEqual([main._tur(x) for x in ("food", "Bar", "İçecek", "x")], ["Yiyecek", "İçecek", "İçecek", None])
        with self.assertRaises(ValueError):
            main.analiz([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "donem.csv", cikti)
            self.assertEqual([(t["tur"], t["maliyet"], t["oran"], t["hedef"]) for t in r["turler"] + [r["genel"]]],
                             [("Yiyecek", D(1081000), D("31.08"), D("31.85")), ("İçecek", D(518400), D("24.26"), D("22.70")), ("Yiyecek + İçecek", D(1599400), D("28.48"), D("28.37"))])
            self.assertEqual([u["kim"] for u in r["uyarilar"] if u["onem"] == "Orta"], ["A la Carte · Yiyecek", "Banket · İçecek", "Havuz Bar · İçecek"])
            self.assertEqual((r["turler"][0]["onceki"], r["turler"][1]["onceki"]), (D("31.56"), None))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Maliyet Oranı", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "d.csv"
            y.write_text("Outlet;Kategori;Açılış;Alımlar;Sayım;Ciro;Hedef\nA;Food;1.000;5.000;800;20.000;30\nA;food;1;1;1;1;\nB;Tatlı;1;1;1;1;\nC;Bar;x;1;1;1;\nD;Bar;1;1;1;1;y\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["dep"], x["tur"], x["acilis"], x["kapanis"]) for x in l], [u["tur"] for u in uy]),
                             ([("A", "Yiyecek", D(1000), D(800))], ["Satır iki kez yazılmış", "Satır okunamadı", "Satır okunamadı", "Satır okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            self.assertEqual(main.main(["--donem", str(y), "--tolerans", "2", "--gun", "31", "--cikti", c]), 0)
            self.assertEqual([main.main(["--donem", str(y), "--cikti", c] + v) for v in (["--tolerans", "x"], ["--gun", "0"])], [2, 2])
            self.assertEqual(main.main(["--donem", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Outlet;Kategori;Sayım;Ciro\nA;x;1;1\n", encoding="utf-8")
            self.assertEqual(main.main(["--donem", str(y), "--cikti", c]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
