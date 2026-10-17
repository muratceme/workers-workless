"""
Yiyecek-İçecek Maliyet Oranı Raporu — Gerisi Bende kod bloğu
Yiyecek ve İçecek › Yiyecek İçecek Müdürü
Dönem stok hareketleri ve satışlardan departman (outlet) bazında yiyecek ve içecek maliyet oranını hesaplar.
  Tüketim maliyeti = açılış stok + alım + transfer giriş − transfer çıkış − kapanış stok
  Satılan maliyet  = tüketim maliyeti − personel yemeği − ikram (complimentary); bu iki kalem başka gider hesabına aktarılır
  Maliyet oranı %  = satılan maliyet ÷ net satış × 100;  fark = oran − hedef (puan)
  Stok günü        = ortalama stok ((açılış + kapanış) ÷ 2) ÷ (tüketim maliyeti ÷ dönem günü)
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
SUTUN = {"dep": ("departman", "outlet", "satis noktasi", "birim"), "tur": ("tur", "kategori", "grup"), "acilis": ("acilis stok", "acilis stogu", "donem basi stok", "acilis"),
         "alim": ("alim", "alimlar", "satin alma", "giris"), "tr_giris": ("transfer giris", "transfer gelen", "gelen transfer"), "tr_cikis": ("transfer cikis", "transfer giden", "giden transfer"),
         "kapanis": ("kapanis stok", "kapanis stogu", "donem sonu stok", "sayim", "kapanis"), "personel": ("personel yemegi", "personel", "calisan yemegi"),
         "ikram": ("ikram", "complimentary", "ikramlar", "comp"), "satis": ("net satis", "satis", "gelir", "ciro"), "hedef": ("hedef", "hedef oran", "hedef maliyet orani"),
         "onceki": ("onceki donem", "onceki oran", "gecen donem")}
SAYI = ("acilis", "alim", "tr_giris", "tr_cikis", "kapanis", "personel", "ikram", "satis")
SIFIR = Decimal(0)


def _tur(x) -> str | None:
    k = katla(x)
    if k in ("yiyecek", "food", "mutfak", "y"):
        return "Yiyecek"
    if k in ("icecek", "beverage", "bar", "i", "alkollu icecek", "alkolsuz icecek"):
        return "İçecek"
    return None


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("dep", "tur", "kapanis", "satis")):
        if not metin(r["dep"]):
            continue
        tur = _tur(r["tur"])
        d = {a: (para(r[a]) if metin(r[a]) else SIFIR) for a in SAYI}
        hedef = para(r["hedef"]) if metin(r["hedef"]) else None
        onceki = para(r["onceki"]) if metin(r["onceki"]) else None
        kim = f"{metin(r['dep'])} · {tur or metin(r['tur'])}"
        if tur is None or any(v is None for v in d.values()) or (metin(r["hedef"]) and hedef is None) or (metin(r["onceki"]) and onceki is None):
            uy.append(uyari("Yüksek", "Satır okunamadı", kim, f"Satır {r['_satir']}: tür ('{metin(r['tur'])}' — Yiyecek / İçecek) ya da tutarlardan biri okunamadı."))
            continue
        if (katla(r["dep"]), tur) in gorulen:
            uy.append(uyari("Orta", "Satır iki kez yazılmış", kim, f"Satır {r['_satir']}: ikinci satır alınmadı."))
            continue
        gorulen.add((katla(r["dep"]), tur))
        l.append({"dep": metin(r["dep"]), "tur": tur, **d, "hedef": hedef, "onceki": onceki})
    return l, uy


def _hesap(s: dict, gun: int) -> dict:
    s["tuketim"] = s["acilis"] + s["alim"] + s["tr_giris"] - s["tr_cikis"] - s["kapanis"]
    s["maliyet"] = s["tuketim"] - s["personel"] - s["ikram"]
    s["oran"] = yuvarla(s["maliyet"] * 100 / s["satis"]) if s["satis"] > 0 else None
    s["fark"] = s["oran"] - s["hedef"] if s["oran"] is not None and s["hedef"] is not None else None
    s["degisim"] = s["oran"] - s["onceki"] if s["oran"] is not None and s["onceki"] is not None else None
    s["stok_gun"] = yuvarla((s["acilis"] + s["kapanis"]) / 2 / (s["tuketim"] / gun), "0.1") if s["tuketim"] > 0 else None
    return s


def analiz(satirlar: list[dict], tolerans: Decimal = Decimal(1), gun: int = 30) -> dict:
    if not satirlar:
        raise ValueError("Okunabilen satır yok.")
    uy = []
    for s in satirlar:
        _hesap(s, gun)
        kim = f"{s['dep']} · {s['tur']}"
        if s["satis"] <= 0:
            uy.append(uyari("Yüksek" if s["maliyet"] > 0 else "Bilgi", "Satış yok", kim, f"Net satış {s['satis']:,.2f}; maliyet oranı hesaplanmadı (satılan maliyet {s['maliyet']:,.2f})."))
        if s["kapanis"] < 0 or s["maliyet"] < 0:
            uy.append(uyari("Yüksek", "Negatif stok / maliyet", kim, f"Kapanış stok {s['kapanis']:,.2f}, satılan maliyet {s['maliyet']:,.2f}; sayım, alım ve transfer kayıtlarını kontrol edin."))
        if s["fark"] is not None and s["fark"] > tolerans:
            uy.append(uyari("Orta", "Maliyet oranı hedefin üstünde", kim, f"Oran %{s['oran']}, hedef %{s['hedef']} (+{s['fark']} puan, tolerans {tolerans}). Fazla maliyet ≈ {yuvarla(s['fark'] * s['satis'] / 100):,.2f}."))
        elif s["fark"] is not None and s["fark"] < -5:
            uy.append(uyari("Bilgi", "Maliyet oranı hedefin çok altında", kim, f"Oran %{s['oran']}, hedef %{s['hedef']} ({s['fark']} puan). Sayım fazlası, eksik alım faturası ya da porsiyon / kalite sorunu olabilir."))
    tr_g = sum((s["tr_giris"] for s in satirlar), SIFIR)
    tr_c = sum((s["tr_cikis"] for s in satirlar), SIFIR)
    if abs(tr_g - tr_c) > 1:
        uy.append(uyari("Bilgi", "Transferler denk değil", f"Giriş {tr_g:,.2f} · çıkış {tr_c:,.2f}", "Departmanlar arası transfer girişleri çıkışlara eşit değil; listede olmayan bir departmana (ör. personel mutfağı, banket) transfer olabilir."))

    def topla(ad, l, tur=""):
        s = {"dep": ad, "tur": tur, **{a: sum((x[a] for x in l), SIFIR) for a in SAYI}}
        hedefli = [x for x in l if x["hedef"] is not None and x["satis"] > 0]
        s["hedef"] = yuvarla(sum((x["hedef"] * x["satis"] for x in hedefli), SIFIR) / sum((x["satis"] for x in hedefli), SIFIR)) if hedefli else None
        onceki = [x for x in l if x["onceki"] is not None and x["satis"] > 0]
        s["onceki"] = None if len(onceki) != len(l) else yuvarla(sum((x["onceki"] * x["satis"] for x in onceki), SIFIR) / sum((x["satis"] for x in onceki), SIFIR))
        return _hesap(s, gun)
    turler = [topla("TOPLAM", [s for s in satirlar if s["tur"] == t], t) for t in ("Yiyecek", "İçecek") if any(s["tur"] == t for s in satirlar)]
    genel = topla("GENEL TOPLAM", satirlar, "Yiyecek + İçecek")
    sira = sorted(satirlar, key=lambda s: (s["tur"] != "Yiyecek", katla(s["dep"])))
    return {"satirlar": sira, "turler": turler, "genel": genel, "uyarilar": uy}


BASLIK = ["Departman", "Tür", "Açılış Stok", "Alım", "Transfer Giriş", "Transfer Çıkış", "Kapanış Stok", "Tüketim Maliyeti", "Personel Yemeği", "İkram", "Satılan Maliyet",
          "Net Satış", "Maliyet %", "Hedef %", "Fark (puan)", "Önceki Dönem %", "Değişim (puan)", "Stok Günü"]


def rapor_yaz(cikti: Path, r: dict, tolerans: Decimal) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Maliyet Oranı"
    _baslik(ws, BASLIK, (20, 14) + (12,) * 16, "C2")
    for s in r["satirlar"] + r["turler"] + [r["genel"]]:
        n = _satir(ws, [s["dep"], s["tur"], s["acilis"], s["alim"], s["tr_giris"], s["tr_cikis"], s["kapanis"], s["tuketim"], s["personel"], s["ikram"], s["maliyet"], s["satis"],
                        s["oran"], s["hedef"], s["fark"], s["onceki"], s["degisim"], s["stok_gun"]], para_s=tuple(range(3, 18)),
                   dolgu={13: KIRMIZI if s["fark"] is not None and s["fark"] > tolerans else (YESIL if s["fark"] is not None else None)})
        if "TOPLAM" in s["dep"]:
            for h in ws[n]:
                h.font = Font(bold=True)
    ws.append([])
    ws.append(["Tüketim maliyeti = açılış + alım + transfer giriş − transfer çıkış − kapanış. Satılan maliyet = tüketim − personel yemeği − ikram. Oran = satılan maliyet ÷ net satış."])
    ws.append(["Toplam satırlarında hedef ve önceki dönem oranı satışla ağırlıklandırılmıştır."])
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, tolerans: Decimal = Decimal(1), gun: int = 30) -> dict:
    satirlar, uy = oku(yol)
    r = analiz(satirlar, tolerans, gun)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, tolerans)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Departman bazında yiyecek ve içecek maliyet oranını hesaplar.")
    p.add_argument("--donem", type=Path, help="Departman, Tür (Yiyecek / İçecek), Açılış Stok, Alım, Transfer Giriş, Transfer Çıkış, Kapanış Stok, Personel Yemeği, İkram, Net Satış, Hedef %%, Önceki Dönem %%")
    p.add_argument("--tolerans", default="1", help="Hedefin kaç puan üstü uyarılır (varsayılan 1)")
    p.add_argument("--gun", type=int, default=30, help="Dönemin gün sayısı, stok günü için (varsayılan 30)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "yiyecek_icecek_maliyet_orani.xlsx")
    a = p.parse_args(argv)
    yol = a.donem or ORNEK / "donem.csv"
    tol = para(a.tolerans)
    if tol is None or not 0 <= tol <= 100 or a.gun < 1:
        print("[X] --tolerans 0-100, --gun en az 1 olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, tol, a.gun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    for s in r["turler"] + [r["genel"]]:
        print(f"[OK] {s['tur']}: satılan maliyet {s['maliyet']:,.2f} · net satış {s['satis']:,.2f} · oran %{s['oran']}" + (f" (hedef %{s['hedef']})" if s["hedef"] is not None else ""))
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
