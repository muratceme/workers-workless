"""
Haftalık Üretim Çizelgesi — Workers / Workless kod bloğu
Üretim › Üretim Planlama › Üretim Planlama Uzmanı

Açık siparişleri, ürün rotalarını (operasyon sırası, iş merkezi, birim süre, hazırlık süresi) ve makine
kapasitelerini (vardiya sayısı, verimlilik) kullanarak sonlu kapasiteli ileri çizelge üretir:
  1. Siparişler en erken termine göre sıralanır (EDD); eşitlikte önceliği yüksek olan önce.
  2. Her siparişin operasyonları rota sırasıyla, kendi iş merkezindeki makineler arasından en erken
     bitirebilecek makineye yerleştirilir; bir operasyon önceki operasyon bitmeden başlamaz.
  3. Süreler çalışma takvimine yayılır: çalışma günleri, 08:00'de başlayan 8 saatlik vardiyalar; net süre
     = (hazırlık + miktar × birim süre) ÷ verimlilik. Vardiya bitince iş sonraki vardiyada devam eder.
Çıktı: sipariş planı (başlangıç, bitiş, termin uyumu), operasyon planı, makine × gün haftalık çizelge,
makine doluluğu ve darboğaz. İnternete bağlanmaz.

Kullanım:
    python main.py                                         # örnek siparişlerle dener
    python main.py --siparisler siparisler.xlsx --rotalar rotalar.xlsx --makineler makineler.xlsx --baslangic 02.11.2026
    python main.py ... --calisma-gunleri 6 --tatiller tatiller.txt
"""
from __future__ import annotations

import argparse
import csv
import math
import sys
from collections import OrderedDict, defaultdict
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
VARDIYA_SAAT = 8
VARDIYA_BASLANGIC = time(8, 0)
GUNLER = ["Pzt", "Sal", "Çar", "Per", "Cum", "Cmt", "Paz"]


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def sayi(x, varsayilan: float = 0.0) -> float:
    if x in (None, ""):
        return varsayilan
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip().replace("%", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(Decimal(s))
    except InvalidOperation:
        return varsayilan


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(str(x).strip(), f).date()
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def sutunlar(baslik, alanlar):
    b = [kucuk(x) for x in baslik]
    return {a: next((i for i, x in enumerate(b) if x in adlar), None) for a, adlar in alanlar.items()}


# ----------------------------------------------------------------------------
# Takvim
# ----------------------------------------------------------------------------

class Takvim:
    """Makine bazında çalışma pencereleri: her çalışma günü 08:00'den başlayan vardiya × 8 saat."""

    def __init__(self, baslangic: date, calisma_gunleri: int, tatiller: set[date]):
        self.baslangic, self.calisma_gunleri, self.tatiller = baslangic, calisma_gunleri, tatiller

    def calisma_gunu(self, g: date) -> bool:
        return g.weekday() < self.calisma_gunleri and g not in self.tatiller

    def pencereler(self, vardiya: int, bas: datetime):
        """bas anından itibaren sıralı çalışma pencereleri (başlangıç, bitiş)."""
        g = min(bas.date(), self.baslangic) if bas.date() < self.baslangic else bas.date() - timedelta(days=1)
        for _ in range(3700):
            if self.calisma_gunu(g):
                p0 = datetime.combine(g, VARDIYA_BASLANGIC)
                p1 = p0 + timedelta(hours=VARDIYA_SAAT * min(vardiya, 3))
                if p1 > bas:
                    yield max(p0, bas), p1
            g += timedelta(days=1)
        raise SystemExit("Çizelge 10 yıllık ufku aştı: kapasite veya rota verisini kontrol edin.")

    def ilerle(self, vardiya: int, bas: datetime, sure_dk: float) -> tuple[datetime, datetime, list[tuple[date, float]]]:
        """sure_dk çalışma dakikasını pencerelere yayar. Dönüş: (gerçek başlangıç, bitiş, [(gün, dakika)])."""
        kalan, ilk, parcalar = sure_dk, None, []
        bas = max(bas, datetime.combine(self.baslangic, VARDIYA_BASLANGIC))
        for p0, p1 in self.pencereler(vardiya, bas):
            if ilk is None:
                ilk = p0
            uygun = (p1 - p0).total_seconds() / 60
            kullan = min(uygun, kalan)
            if kullan > 0:
                parcalar.append((p0.date() if p0.time() >= VARDIYA_BASLANGIC else p0.date() - timedelta(days=1), kullan))
            kalan -= kullan
            if kalan <= 1e-9:
                return ilk, p0 + timedelta(minutes=kullan), parcalar
        raise AssertionError("ulaşılamaz")


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

def verileri_oku(siparis_yolu: Path, rota_yolu: Path, makine_yolu: Path):
    s = tablo_oku(siparis_yolu)
    k = sutunlar(s[0], {"no": ("sipariş no", "iş emri", "iş emri no", "no"), "urun": ("ürün", "ürün kodu", "stok kodu", "mamul"),
                        "miktar": ("miktar", "adet"), "termin": ("termin", "termin tarihi", "teslim tarihi"),
                        "oncelik": ("öncelik",), "musteri": ("müşteri",)})
    if None in (k["no"], k["urun"], k["miktar"], k["termin"]):
        raise SystemExit(f"Siparişlerde Sipariş No, Ürün, Miktar ve Termin gerekli. Başlıklar: {s[0]}")
    siparisler = []
    for r in s[1:]:
        if r[k["no"]] in (None, ""):
            continue
        siparisler.append({"no": str(r[k["no"]]).strip(), "urun": str(r[k["urun"]]).strip(), "miktar": sayi(r[k["miktar"]]),
                           "termin": tarih(r[k["termin"]]), "oncelik": sayi(r[k["oncelik"]], 0) if k["oncelik"] is not None else 0,
                           "musteri": str(r[k["musteri"]] or "") if k["musteri"] is not None else ""})

    s = tablo_oku(rota_yolu)
    k = sutunlar(s[0], {"urun": ("ürün", "ürün kodu", "stok kodu", "mamul"), "sira": ("operasyon sırası", "sıra", "op no"),
                        "op": ("operasyon", "operasyon adı"), "grup": ("iş merkezi", "makine grubu", "istasyon"),
                        "birim": ("birim süre (dk)", "birim süre", "çevrim süresi (dk)"), "hazirlik": ("hazırlık süresi (dk)", "hazırlık süresi", "setup (dk)")})
    if None in (k["urun"], k["grup"], k["birim"]):
        raise SystemExit(f"Rotalarda Ürün, İş Merkezi ve Birim Süre gerekli. Başlıklar: {s[0]}")
    rotalar: dict[str, list[dict]] = defaultdict(list)
    for i, r in enumerate(s[1:]):
        if r[k["urun"]] in (None, ""):
            continue
        rotalar[str(r[k["urun"]]).strip()].append({
            "sira": sayi(r[k["sira"]], i) if k["sira"] is not None else i, "op": str(r[k["op"]] or "") if k["op"] is not None else "",
            "grup": str(r[k["grup"]]).strip(), "birim": sayi(r[k["birim"]]), "hazirlik": sayi(r[k["hazirlik"]]) if k["hazirlik"] is not None else 0.0})
    for v in rotalar.values():
        v.sort(key=lambda x: x["sira"])

    s = tablo_oku(makine_yolu)
    k = sutunlar(s[0], {"makine": ("makine", "makine kodu", "tezgâh", "tezgah"), "grup": ("iş merkezi", "makine grubu", "istasyon"),
                        "vardiya": ("vardiya", "vardiya sayısı"), "verim": ("verimlilik", "verimlilik %", "oee")})
    if None in (k["makine"], k["grup"]):
        raise SystemExit(f"Makinelerde Makine ve İş Merkezi gerekli. Başlıklar: {s[0]}")
    makineler = OrderedDict()
    for r in s[1:]:
        if r[k["makine"]] in (None, ""):
            continue
        v = sayi(r[k["verim"]], 100) if k["verim"] is not None else 100
        makineler[str(r[k["makine"]]).strip()] = {"grup": str(r[k["grup"]]).strip(), "vardiya": int(sayi(r[k["vardiya"]], 1)) if k["vardiya"] is not None else 1,
                                                  "verim": v / 100 if v > 1 else v}
    return siparisler, rotalar, makineler


# ----------------------------------------------------------------------------
# Çizelgeleme
# ----------------------------------------------------------------------------

def cizelgele(siparisler, rotalar, makineler, takvim: Takvim) -> dict:
    serbest = {m: datetime.combine(takvim.baslangic, VARDIYA_BASLANGIC) for m in makineler}
    yuk: dict[str, dict[date, float]] = defaultdict(lambda: defaultdict(float))
    operasyonlar, uyarilar, plan = [], [], []
    sira = sorted(siparisler, key=lambda s: (s["termin"] or date.max, -s["oncelik"], s["no"]))
    for sp in sira:
        rota = rotalar.get(sp["urun"])
        if not rota:
            uyarilar.append(f"{sp['no']}: '{sp['urun']}' için rota yok; çizelgelenmedi")
            continue
        hazir = datetime.combine(takvim.baslangic, VARDIYA_BASLANGIC)
        ilk_bas, eksik = None, False
        for op in rota:
            adaylar = [m for m, v in makineler.items() if v["grup"] == op["grup"]]
            if not adaylar:
                uyarilar.append(f"{sp['no']}: '{op['grup']}' iş merkezinde makine yok")
                eksik = True
                break
            en_iyi = None
            for m in adaylar:
                v = makineler[m]
                sure = (op["hazirlik"] + sp["miktar"] * op["birim"]) / v["verim"]
                bas, bit, parca = takvim.ilerle(v["vardiya"], max(serbest[m], hazir), sure)
                if en_iyi is None or bit < en_iyi[2]:
                    en_iyi = (m, bas, bit, parca, sure)
            m, bas, bit, parca, sure = en_iyi
            serbest[m] = bit
            hazir = bit
            for g, dk in parca:
                yuk[m][g] += dk
            ilk_bas = ilk_bas or bas
            operasyonlar.append({"siparis": sp["no"], "urun": sp["urun"], "op": op["op"] or op["grup"], "grup": op["grup"], "makine": m,
                                 "baslangic": bas, "bitis": bit, "sure_dk": sure, "parcalar": parca})
        if eksik:
            continue
        gec = (hazir.date() - sp["termin"]).days if sp["termin"] else None
        plan.append({**sp, "baslangic": ilk_bas, "bitis": hazir, "gecikme": gec if gec and gec > 0 else 0})
    return {"plan": plan, "operasyonlar": operasyonlar, "yuk": yuk, "uyarilar": uyarilar}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
YESIL = PatternFill("solid", fgColor="E3F5E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, b, g):
    ws.append(b)
    for c in ws[ws.max_row]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI
    for j, w in enumerate(g, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "B2"


def calistir(siparis_yolu: Path, rota_yolu: Path, makine_yolu: Path, cikti: Path, baslangic: date, calisma_gunleri: int = 5,
             tatiller: set[date] | None = None, hafta: int = 1) -> dict:
    siparisler, rotalar, makineler = verileri_oku(siparis_yolu, rota_yolu, makine_yolu)
    takvim = Takvim(baslangic, calisma_gunleri, tatiller or set())
    c = cizelgele(siparisler, rotalar, makineler, takvim)
    gunler = [baslangic + timedelta(days=i) for i in range(7 * hafta) if takvim.calisma_gunu(baslangic + timedelta(days=i))]
    kapasite = {m: VARDIYA_SAAT * 60 * min(v["vardiya"], 3) for m, v in makineler.items()}
    doluluk = {m: sum(c["yuk"][m].get(g, 0) for g in gunler) / (kapasite[m] * len(gunler)) if gunler else 0 for m in makineler}
    c["doluluk"], c["gunler"] = doluluk, gunler

    wb = Workbook()
    p = wb.active
    p.title = "Sipariş Planı"
    _baslik(p, ["Sıra", "Sipariş No", "Müşteri", "Ürün", "Miktar", "Termin", "Planlanan Başlangıç", "Planlanan Bitiş", "Gecikme (gün)", "Durum"],
            (6, 12, 18, 14, 9, 12, 18, 18, 12, 14))
    for i, x in enumerate(c["plan"], 1):
        durum = "GECİKECEK" if x["gecikme"] else "Zamanında"
        p.append([i, x["no"], x["musteri"], x["urun"], x["miktar"], x["termin"], x["baslangic"], x["bitis"], x["gecikme"] or None, durum])
        p.cell(p.max_row, 6).number_format = "DD.MM.YYYY"
        p.cell(p.max_row, 7).number_format = p.cell(p.max_row, 8).number_format = "DD.MM.YYYY HH:MM"
        p.cell(p.max_row, 10).fill = KIRMIZI if x["gecikme"] else YESIL
    for u in c["uyarilar"]:
        p.append(["", "Uyarı", u])

    h = wb.create_sheet("Haftalık Çizelge")
    _baslik(h, ["Makine", "İş Merkezi"] + [f"{GUNLER[g.weekday()]} {g:%d.%m}" for g in gunler] + ["Doluluk"], [12, 14] + [26] * len(gunler) + [10])
    gun_is = defaultdict(lambda: defaultdict(list))
    for o in c["operasyonlar"]:
        for g, dk in o["parcalar"]:
            gun_is[o["makine"]][g].append(f"{o['siparis']} {o['op']} ({dk:.0f} dk)")
    for m, v in makineler.items():
        h.append([m, v["grup"]] + ["\n".join(gun_is[m].get(g, [])) or "—" for g in gunler] + [doluluk[m]])
        h.cell(h.max_row, 3 + len(gunler)).number_format = "0%"
        for j, g in enumerate(gunler, 3):
            oran = c["yuk"][m].get(g, 0) / kapasite[m]
            h.cell(h.max_row, j).fill = KIRMIZI if oran >= 0.95 else SARI if oran >= 0.6 else PatternFill()
        for x in h[h.max_row]:
            x.alignment = UST

    o = wb.create_sheet("Operasyon Planı")
    _baslik(o, ["Sipariş No", "Ürün", "Operasyon", "İş Merkezi", "Makine", "Başlangıç", "Bitiş", "Süre (dk)"], (12, 14, 18, 14, 12, 18, 18, 10))
    for x in sorted(c["operasyonlar"], key=lambda x: (x["makine"], x["baslangic"])):
        o.append([x["siparis"], x["urun"], x["op"], x["grup"], x["makine"], x["baslangic"], x["bitis"], round(x["sure_dk"])])
        o.cell(o.max_row, 6).number_format = o.cell(o.max_row, 7).number_format = "DD.MM.YYYY HH:MM"
    o.auto_filter.ref = o.dimensions

    d = wb.create_sheet("Makine Doluluğu")
    _baslik(d, ["Makine", "İş Merkezi", "Vardiya", "Verimlilik", "Hafta Yükü (dk)", "Hafta Kapasitesi (dk)", "Doluluk"], (12, 14, 9, 11, 15, 18, 10))
    for m, v in sorted(makineler.items(), key=lambda i: -doluluk[i[0]]):
        d.append([m, v["grup"], v["vardiya"], v["verim"], round(sum(c["yuk"][m].get(g, 0) for g in gunler)), kapasite[m] * len(gunler), doluluk[m]])
        d.cell(d.max_row, 4).number_format = d.cell(d.max_row, 7).number_format = "0%"
    if makineler:
        gr = BarChart()
        gr.title, gr.height, gr.width = "Makine doluluğu (darboğaz en üstte)", 8, 18
        gr.add_data(Reference(d, min_col=7, min_row=1, max_row=1 + len(makineler)), titles_from_data=True)
        gr.set_categories(Reference(d, min_col=1, min_row=2, max_row=1 + len(makineler)))
        d.add_chart(gr, "I2")

    b = wb.create_sheet("Bilgi")
    for s in [["Sıralama", "En erken termin önce (EDD); eşitlikte öncelik değeri yüksek olan önce"],
              ["Yerleşim", "Her operasyon, iş merkezindeki makineler arasından en erken bitirebilecek makineye; önceki operasyon bitmeden başlamaz"],
              ["Süre", "(Hazırlık + miktar × birim süre) ÷ verimlilik; vardiya bitince iş sonraki vardiyada devam eder"],
              ["Takvim", f"Haftada {calisma_gunleri} çalışma günü, vardiyalar {VARDIYA_BASLANGIC:%H:%M}'de başlar ve {VARDIYA_SAAT} saattir; "
                         "tatil günleri atlanır"],
              ["Not", "Sezgisel bir çizelgedir (optimum değildir); malzeme hazır olma ve kalıp/aparat kısıtları dikkate alınmaz"]]:
        b.append(s)
    b.column_dimensions["A"].width = 14
    b.column_dimensions["B"].width = 120
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    return c


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description="Sonlu kapasiteli haftalık üretim çizelgesi (EDD + ileri çizelgeleme).")
    ap.add_argument("--siparisler", type=Path, default=BURASI / "ornek_veri" / "siparisler.csv",
                    help="Siparişler (.xlsx/.csv): Sipariş No, Ürün, Miktar, Termin; Öncelik, Müşteri isteğe bağlı")
    ap.add_argument("--rotalar", type=Path, default=BURASI / "ornek_veri" / "rotalar.csv",
                    help="Rotalar: Ürün, Operasyon Sırası, Operasyon, İş Merkezi, Birim Süre (dk), Hazırlık Süresi (dk)")
    ap.add_argument("--makineler", type=Path, default=BURASI / "ornek_veri" / "makineler.csv",
                    help="Makineler: Makine, İş Merkezi, Vardiya, Verimlilik %%")
    ap.add_argument("--baslangic", help="Çizelge başlangıç günü (GG.AA.YYYY); varsayılan önümüzdeki pazartesi")
    ap.add_argument("--calisma-gunleri", type=int, default=5, choices=range(1, 8), metavar="1-7", help="Haftalık çalışma günü (5: Pzt-Cum)")
    ap.add_argument("--tatiller", type=Path, help="Tatil günleri (.txt, GG.AA.YYYY)")
    ap.add_argument("--hafta", type=int, default=1, help="Haftalık çizelgede gösterilecek hafta sayısı")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "uretim_cizelgesi.xlsx")
    a = ap.parse_args(argv)
    if a.baslangic:
        bas = tarih(a.baslangic)
    elif a.siparisler == BURASI / "ornek_veri" / "siparisler.csv":
        bas = date(2026, 11, 2)
    else:
        bugun = date.today()
        bas = bugun + timedelta(days=(7 - bugun.weekday()) % 7 or 7)
    tatil = set()
    if a.tatiller:
        tatil = {tarih(x) for x in a.tatiller.read_text(encoding="utf-8").split() if tarih(x)}
    c = calistir(a.siparisler, a.rotalar, a.makineler, a.cikti, bas, a.calisma_gunleri, tatil, a.hafta)
    gec = [x for x in c["plan"] if x["gecikme"]]
    darbogaz = max(c["doluluk"].items(), key=lambda i: i[1]) if c["doluluk"] else None
    print(f"[OK] {len(c['plan'])} sipariş çizelgelendi · başlangıç {bas:%d.%m.%Y} · gecikecek: {len(gec)}")
    for x in gec:
        print(f"[!] {x['no']} ({x['urun']}): termin {x['termin']:%d.%m.%Y}, bitiş {x['bitis']:%d.%m.%Y %H:%M} → {x['gecikme']} gün gecikme")
    if darbogaz:
        print(f"[i] Darboğaz: {darbogaz[0]} (%{darbogaz[1] * 100:.0f} doluluk)")
    for u in c["uyarilar"]:
        print(f"[!] {u}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
