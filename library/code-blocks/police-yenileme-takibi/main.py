"""
Poliçe Yenileme Takibi — Workers / Workless kod bloğu
Teknik ve Operasyon › Sigorta Uzmanı

Vadesi yaklaşan poliçeleri listeler ve müşteri başına tek bir yenileme hatırlatma mesajı hazırlar.
  Kalan gün  = bitiş tarihi − rapor tarihi
  Kapsam     = −(--gecmis-gun, varsayılan 15) ≤ kalan gün ≤ --gun (varsayılan 30);  "Yenilendi" sütunu Evet olanlar alınmaz
  Grup       = Süresi geçmiş (kalan gün < 0) · Bu hafta (0-7) · Yaklaşan (8 ve üzeri)
  Mesaj      = aynı müşterinin kapsamdaki tüm poliçeleri tek mesajda; şablon --sablon ile değiştirilebilir ({musteri}, {policeler}, {imza})
Kod mesaj göndermez; yalnız metni hazırlar. İletişim izni "Hayır" olan müşteriler işaretlenir, gönderim kararı sizindir.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"no": ("police no", "police numarasi"), "musteri": ("musteri", "sigortali", "musteri adi", "ad soyad"), "telefon": ("telefon", "cep telefonu", "gsm"),
         "eposta": ("e posta", "eposta", "email", "e mail"), "brans": ("brans", "urun"), "sirket": ("sigorta sirketi", "sirket"), "konu": ("konu", "plaka", "aciklama", "risk"),
         "bitis": ("bitis tarihi", "bitis", "vade", "vade tarihi"), "prim": ("prim", "son prim", "mevcut prim"), "temsilci": ("musteri temsilcisi", "temsilci", "sorumlu"),
         "izin": ("iletisim izni", "izin", "iys izni", "ileti izni"), "yenilendi": ("yenilendi", "yenileme durumu")}
SABLON = ("Sayın {musteri},\n\n{policeler}\n\nTeminatınızın kesintiye uğramaması için yenileme işlemlerinizi birlikte tamamlayabiliriz. "
          "Uygun olduğunuz bir zamanda bize ulaşmanızı rica ederiz.\n\n{imza}")
EVET, HAYIR = ("evet", "e", "var", "yenilendi", "1", "x"), ("hayir", "h", "yok", "0")
Z = Decimal(0)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("no", "musteri", "bitis")):
        no = metin(r["no"])
        if not no:
            continue
        kim = f"{no} (satır {r['_satir']})"
        if no in gorulen:
            uy.append(uyari("Yüksek", "Mükerrer poliçe", kim, "Poliçe numarası tekrar ediyor; ilk satır kullanıldı."))
            continue
        gorulen.add(no)
        bitis = tarih(r["bitis"])
        if not bitis:
            uy.append(uyari("Yüksek", "Bitiş tarihi okunamadı", kim, f"'{metin(r['bitis'])}'. Poliçe takibe alınamadı."))
            continue
        if not metin(r["musteri"]):
            uy.append(uyari("Yüksek", "Müşteri adı boş", kim, "Poliçe takibe alınamadı."))
            continue
        prim = para(r["prim"])
        if prim is None and metin(r["prim"]):
            uy.append(uyari("Orta", "Prim okunamadı", kim, f"'{metin(r['prim'])}'. Prim boş sayıldı."))
        l.append({"no": no, "musteri": metin(r["musteri"]), "telefon": metin(r["telefon"]), "eposta": metin(r["eposta"]), "brans": metin(r["brans"]) or "Sigorta", "sirket": metin(r["sirket"]),
                  "konu": metin(r["konu"]), "bitis": bitis, "prim": prim, "temsilci": metin(r["temsilci"]) or "Belirtilmemiş", "izin": None if not metin(r["izin"]) else katla(r["izin"]) not in HAYIR,
                  "yenilendi": katla(r["yenilendi"]) in EVET})
    return l, uy


def satir_metni(p: dict) -> str:
    ek = ", ".join(x for x in (p["no"], p["konu"]) if x)
    bas = f"- {p['brans']} poliçeniz ({ek})" + (f", {p['sirket']}" if p["sirket"] else "") + ": "
    g = p["kalan"]
    if g < 0:
        return bas + f"{p['bitis']:%d.%m.%Y} tarihinde sona erdi ({-g} gün önce)."
    if g == 0:
        return bas + "bugün sona eriyor."
    return bas + f"{p['bitis']:%d.%m.%Y} tarihinde sona eriyor ({g} gün kaldı)."


def takip(policeler: list[dict], rapor: date, gun: int = 30, gecmis_gun: int = 15, sablon: str = SABLON, imza: str = "") -> dict:
    kapsam, sayac = [], {"yenilenmis": 0, "ileride": 0, "eski": 0}
    for p in policeler:
        p["kalan"] = (p["bitis"] - rapor).days
        if p["yenilendi"]:
            sayac["yenilenmis"] += 1
        elif p["kalan"] > gun:
            sayac["ileride"] += 1
        elif p["kalan"] < -gecmis_gun:
            sayac["eski"] += 1
        else:
            p["grup"] = "Süresi geçmiş" if p["kalan"] < 0 else "Bu hafta" if p["kalan"] <= 7 else "Yaklaşan"
            kapsam.append(p)
    kapsam.sort(key=lambda p: (p["kalan"], p["musteri"], p["no"]))
    musteri = {}
    for p in kapsam:
        m = musteri.setdefault((katla(p["musteri"]), katla(p["telefon"])), {"musteri": p["musteri"], "telefon": p["telefon"], "eposta": p["eposta"], "temsilci": p["temsilci"], "policeler": [],
                                                                         "izin": p["izin"]})
        m["policeler"].append(p)
        m["telefon"], m["eposta"] = m["telefon"] or p["telefon"], m["eposta"] or p["eposta"]
        if p["izin"] is False:
            m["izin"] = False
    mesajlar = []
    for m in musteri.values():
        m["prim"] = sum((p["prim"] or Z for p in m["policeler"]), Z)
        m["kalan"] = min(p["kalan"] for p in m["policeler"])
        m["mesaj"] = sablon.format(musteri=m["musteri"], policeler="\n".join(satir_metni(p) for p in m["policeler"]), imza=imza or m["temsilci"]).strip()
        m["not"] = "İletişim izni 'Hayır': gönderim öncesi değerlendirin." if m["izin"] is False else "" if m["telefon"] or m["eposta"] else "Telefon ve e-posta yok."
        mesajlar.append(m)
    temsilci = {}
    for p in kapsam:
        t = temsilci.setdefault(p["temsilci"], {"ad": p["temsilci"], "police": 0, "prim": Z, "Süresi geçmiş": 0, "Bu hafta": 0, "Yaklaşan": 0})
        t["police"] += 1
        t["prim"] += p["prim"] or Z
        t[p["grup"]] += 1
    return {"kapsam": kapsam, "mesajlar": mesajlar, "temsilci": sorted(temsilci.values(), key=lambda t: (-t["police"], t["ad"])), "sayac": sayac,
            "prim": sum((p["prim"] or Z for p in kapsam), Z), "gun": gun, "gecmis_gun": gecmis_gun}


def mesaj_dosyasi(r: dict) -> str:
    if not r["mesajlar"]:
        return "Kapsamda yenilemesi yaklaşan poliçe yok.\n"
    b = []
    for m in r["mesajlar"]:
        b.append(f"=== {m['musteri']} · {m['telefon'] or 'telefon yok'} · {m['eposta'] or 'e-posta yok'} ===" + (f"\n[!] {m['not']}" if m["not"] else "") + f"\n\n{m['mesaj']}\n")
    return "\n".join(b)


def rapor_yaz(cikti: Path, r: dict, rapor: date, uy: list[dict]) -> None:
    wb = Workbook()
    renk = {"Süresi geçmiş": KIRMIZI, "Bu hafta": SARI, "Yaklaşan": None}
    ws = wb.active
    ws.title = "Yenileme Listesi"
    s = r["sayac"]
    ws.append([f"Rapor tarihi {rapor:%d.%m.%Y} · kapsam: {r['gecmis_gun']} gün önce biten – {r['gun']} gün içinde bitecek · {len(r['kapsam'])} poliçe, prim {r['prim']:,.2f} · "
               f"alınmayan: {s['yenilenmis']} yenilenmiş, {s['ileride']} ileri tarihli, {s['eski']} daha eski"])
    ws["A1"].font = Font(bold=True)
    _baslik(ws, ["Grup", "Kalan Gün", "Bitiş Tarihi", "Poliçe No", "Müşteri", "Branş", "Sigorta Şirketi", "Konu", "Prim", "Telefon", "E-posta", "Müşteri Temsilcisi", "İletişim İzni",
                 "Arandı / Sonuç (siz doldurun)"], (15, 9, 12, 14, 30, 14, 20, 24, 14, 16, 26, 20, 11, 30), "A3")
    for p in r["kapsam"]:
        _satir(ws, [p["grup"], p["kalan"], p["bitis"], p["no"], p["musteri"], p["brans"], p["sirket"], p["konu"], p["prim"], p["telefon"], p["eposta"], p["temsilci"],
                    "" if p["izin"] is None else "Evet" if p["izin"] else "Hayır", ""], para_s=(9,), tarih_s=(3,), dolgu={1: renk[p["grup"]], 13: KIRMIZI if p["izin"] is False else None})
    ws = wb.create_sheet("Mesajlar")
    _baslik(ws, ["Müşteri", "Telefon", "E-posta", "Poliçe Sayısı", "Toplam Prim", "En Yakın Bitiş (gün)", "Not", "Mesaj Taslağı", "Gönderildi (siz doldurun)"], (30, 16, 26, 9, 14, 10, 34, 100, 14))
    for m in r["mesajlar"]:
        _satir(ws, [m["musteri"], m["telefon"], m["eposta"], len(m["policeler"]), m["prim"], m["kalan"], m["not"], m["mesaj"], ""], para_s=(5,), dolgu={7: SARI if m["not"] else None}, sar=(7, 8))
    ws = wb.create_sheet("Temsilci Özeti")
    _baslik(ws, ["Müşteri Temsilcisi", "Poliçe", "Süresi Geçmiş", "Bu Hafta", "Yaklaşan", "Toplam Prim"], (26, 9, 12, 10, 10, 16))
    for t in r["temsilci"]:
        _satir(ws, [t["ad"], t["police"], t["Süresi geçmiş"] or None, t["Bu hafta"] or None, t["Yaklaşan"] or None, t["prim"]], para_s=(6,), dolgu={3: KIRMIZI if t["Süresi geçmiş"] else None})
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)
    (cikti.parent / "yenileme_mesajlari.txt").write_text(mesaj_dosyasi(r), encoding="utf-8")


def calistir(yol: Path, cikti: Path, rapor: date, **secenek) -> dict:
    policeler, uy = oku(yol)
    if not policeler:
        raise ValueError(f"{yol.name}: takip edilecek poliçe yok.")
    r = takip(policeler, rapor, **secenek)
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, rapor, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Vadesi yaklaşan poliçeleri listeler ve müşteri başına yenileme hatırlatma mesajı hazırlar.")
    p.add_argument("--policeler", type=Path, default=ORNEK / "policeler.csv",
                   help="Poliçe No, Müşteri, Telefon, E-posta, Branş, Sigorta Şirketi, Konu, Bitiş Tarihi, Prim, Müşteri Temsilcisi, İletişim İzni, Yenilendi")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 01.10.2026)")
    p.add_argument("--gun", type=int, default=30, help="Kaç gün içinde bitecek poliçeler alınsın (varsayılan 30)")
    p.add_argument("--gecmis-gun", type=int, default=15, help="Süresi en çok kaç gün önce bitmiş poliçeler alınsın (varsayılan 15)")
    p.add_argument("--sablon", type=Path, help="Mesaj şablonu (.txt): {musteri}, {policeler}, {imza} alanlarını içerir")
    p.add_argument("--imza", default="", help="Mesaj sonundaki imza (varsayılan: müşteri temsilcisinin adı)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "yenileme_takibi.xlsx")
    a = p.parse_args(argv)
    for y in (a.policeler, a.sablon):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    rapor = tarih(a.tarih) if a.tarih else (date(2026, 10, 1) if a.policeler == ORNEK / "policeler.csv" else date.today())
    sablon = a.sablon.read_text(encoding="utf-8-sig") if a.sablon else SABLON
    try:
        sablon.format(musteri="", policeler="", imza="")
        gecerli = "{policeler}" in sablon
    except (KeyError, IndexError, ValueError):
        gecerli = False
    if not rapor or a.gun < 0 or a.gecmis_gun < 0 or not gecerli:
        print("[X] --tarih GG.AA.YYYY biçiminde; --gun ve --gecmis-gun sıfır veya pozitif; şablon {policeler} içermeli ve yalnız {musteri}, {policeler}, {imza} alanlarını kullanmalı")
        return 2
    try:
        r = calistir(a.policeler, a.cikti, rapor, gun=a.gun, gecmis_gun=a.gecmis_gun, sablon=sablon, imza=a.imza)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    say = lambda g: sum(x["grup"] == g for x in r["kapsam"])   # noqa: E731
    print(f"[OK] {len(r['kapsam'])} poliçe kapsamda · {say('Süresi geçmiş')} süresi geçmiş · {say('Bu hafta')} bu hafta · {say('Yaklaşan')} yaklaşan · {len(r['mesajlar'])} mesaj taslağı")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[OK] Mesajlar: {(a.cikti.parent / 'yenileme_mesajlari.txt').resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
