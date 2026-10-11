import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
RAPOR = date(2026, 10, 1)


def pol(no, gun, musteri="Ayşe Örnek", telefon="0500", prim=1000, izin=True, yenilendi=False, temsilci="T1", brans="Kasko", konu="34 ABC 001", sirket="X Sigorta"):
    return {"no": no, "musteri": musteri, "telefon": telefon, "eposta": "", "brans": brans, "sirket": sirket, "konu": konu, "bitis": RAPOR + timedelta(days=gun),
            "prim": None if prim is None else D(prim), "temsilci": temsilci, "izin": izin, "yenilendi": yenilendi}


class TakipTesti(unittest.TestCase):
    def test_kapsam_ve_gruplar(self):
        l = [pol(f"P{g}", g, musteri=f"M{g}") for g in (-16, -15, -1, 0, 7, 8, 30, 31)] + [pol("PY", 5, yenilendi=True)]
        r = main.takip(l, RAPOR)
        self.assertEqual([(p["no"], p["grup"]) for p in r["kapsam"]],
                         [("P-15", "Süresi geçmiş"), ("P-1", "Süresi geçmiş"), ("P0", "Bu hafta"), ("P7", "Bu hafta"), ("P8", "Yaklaşan"), ("P30", "Yaklaşan")])
        self.assertEqual((r["sayac"], r["prim"]), ({"yenilenmis": 1, "ileride": 1, "eski": 1}, D(6000)))
        self.assertEqual([p["no"] for p in main.takip(l, RAPOR, gun=7, gecmis_gun=0)["kapsam"]], ["P0", "P7"])

    def test_mesaj_metni(self):
        r = main.takip([pol("P1", 4), pol("P2", 4, brans="Trafik", prim=500), pol("P3", -3, musteri="Firma A.Ş.", telefon="0501", konu="", sirket=""), pol("P4", 0, musteri="Bugün", telefon="")],
                       RAPOR, imza="Örnek Sigorta Aracılık")
        m = {x["musteri"]: x for x in r["mesajlar"]}
        self.assertEqual((len(r["mesajlar"]), len(m["Ayşe Örnek"]["policeler"]), m["Ayşe Örnek"]["prim"], m["Ayşe Örnek"]["kalan"]), (3, 2, D(1500), 4))
        self.assertIn("Sayın Ayşe Örnek,\n\n- Kasko poliçeniz (P1, 34 ABC 001), X Sigorta: 05.10.2026 tarihinde sona eriyor (4 gün kaldı).\n- Trafik poliçeniz (P2", m["Ayşe Örnek"]["mesaj"])
        self.assertTrue(m["Ayşe Örnek"]["mesaj"].endswith("Örnek Sigorta Aracılık"))
        self.assertIn("- Kasko poliçeniz (P3): 28.09.2026 tarihinde sona erdi (3 gün önce).", m["Firma A.Ş."]["mesaj"])
        self.assertIn("bugün sona eriyor.", m["Bugün"]["mesaj"])
        self.assertEqual((m["Bugün"]["not"], m["Ayşe Örnek"]["not"]), ("Telefon ve e-posta yok.", ""))
        self.assertTrue(main.takip([pol("P1", 4)], RAPOR)["mesajlar"][0]["mesaj"].endswith("T1"))                 # imza verilmezse temsilci

    def test_izin_sablon_ve_adaslar(self):
        r = main.takip([pol("P1", 4, izin=True), pol("P2", 9, izin=False), pol("P3", 5, telefon="0599")], RAPOR, sablon="{musteri} | {policeler} | {imza}")
        self.assertEqual([(len(m["policeler"]), m["izin"], m["telefon"]) for m in r["mesajlar"]], [(2, False, "0500"), (1, True, "0599")])   # aynı ad, farklı telefon ayrı müşteri
        self.assertIn("İletişim izni 'Hayır'", r["mesajlar"][0]["not"])
        self.assertTrue(r["mesajlar"][1]["mesaj"].startswith("Ayşe Örnek | - Kasko poliçeniz (P3"))
        self.assertIn("[!] İletişim izni", main.mesaj_dosyasi(r))
        self.assertEqual(main.mesaj_dosyasi(main.takip([pol("P1", 90)], RAPOR)), "Kapsamda yenilemesi yaklaşan poliçe yok.\n")

    def test_temsilci_ozeti(self):
        r = main.takip([pol("P1", -2, temsilci="A"), pol("P2", 3, temsilci="A", prim=None), pol("P3", 20, temsilci="B"), pol("P4", 21, temsilci="A")], RAPOR)
        self.assertEqual([(t["ad"], t["police"], t["Süresi geçmiş"], t["Bu hafta"], t["Yaklaşan"], t["prim"]) for t in r["temsilci"]], [("A", 3, 1, 1, 1, D(2000)), ("B", 1, 0, 0, 1, D(1000))])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "y.xlsx"
            r = main.calistir(ORNEK / "policeler.csv", cikti, RAPOR)
            self.assertEqual([(p["no"], p["kalan"], p["grup"]) for p in r["kapsam"]],
                             [("Y-04", -3, "Süresi geçmiş"), ("Y-07", 0, "Bu hafta"), ("Y-01", 4, "Bu hafta"), ("Y-02", 4, "Bu hafta"), ("Y-03", 19, "Yaklaşan"), ("Y-05", 30, "Yaklaşan")])
            self.assertEqual((r["prim"], len(r["mesajlar"]), r["sayac"]), (D("283700.00"), 5, {"yenilenmis": 1, "ileride": 1, "eski": 1}))
            self.assertEqual([m["musteri"] for m in r["mesajlar"] if m["not"]], ["Mehmet Deneme"])
            self.assertEqual([u["tur"] for u in r["uyarilar"]], ["Bitiş tarihi okunamadı"])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Yenileme Listesi", "Mesajlar", "Temsilci Özeti", "Uyarılar"])
            t = (Path(tmp) / "yenileme_mesajlari.txt").read_text(encoding="utf-8")
            self.assertIn("- Trafik poliçeniz (Y-02, 34 ABC 001), Kurgu Sigorta: 05.10.2026 tarihinde sona eriyor (4 gün kaldı).", t)

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, s = Path(tmp) / "p.csv", Path(tmp) / "s.txt"
            y.write_text("Poliçe No;Müşteri;Bitiş Tarihi;Prim;İletişim İzni;Yenilendi\nP1;A;05.10.2026;çok;HAYIR;\nP1;A;06.10.2026;;;\nP2;;05.10.2026;;;\nP3;B;yakında;;;\nP4;C;07.10.2026;10;;evet\n",
                         encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual([(p["no"], p["prim"], p["izin"], p["yenilendi"], p["brans"]) for p in l], [("P1", None, False, False, "Sigorta"), ("P4", D(10), None, True, "Sigorta")])
            self.assertEqual(sorted(u["tur"] for u in uy), ["Bitiş tarihi okunamadı", "Mükerrer poliçe", "Müşteri adı boş", "Prim okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--policeler", str(y), "--cikti", c, "--tarih", "01.10.2026"]
            s.write_text("Merhaba {musteri}\n{policeler}\n{imza}", encoding="utf-8")
            self.assertEqual(main.main(arg + ["--sablon", str(s), "--imza", "Acente", "--gun", "10"]), 0)
            self.assertIn("Merhaba A\n- Sigorta poliçeniz (P1)", (Path(tmp) / "yenileme_mesajlari.txt").read_text(encoding="utf-8"))
            s.write_text("Merhaba {ad} {policeler}", encoding="utf-8")
            self.assertEqual(main.main(arg + ["--sablon", str(s)]), 2)
            s.write_text("Merhaba {musteri}", encoding="utf-8")
            self.assertEqual(main.main(arg + ["--sablon", str(s)]), 2)
            self.assertEqual(main.main(arg[:-1] + ["dün"]), 2)
            self.assertEqual(main.main(["--policeler", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Poliçe No;Müşteri;Bitiş Tarihi\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
