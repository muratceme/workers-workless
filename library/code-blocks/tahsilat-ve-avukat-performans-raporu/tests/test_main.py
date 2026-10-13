import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)


def d(no, t=date(2026, 5, 1), tutar=1000, tahsilat=0, masraf=0, avukat="A", durum="Açık"):
    return {"no": no, "avukat": avukat, "tarih": t, "tutar": D(str(tutar)), "tahsilat": D(str(tahsilat)), "masraf": D(str(masraf)), "durum": durum}


class RaporTesti(unittest.TestCase):
    def test_yas(self):
        self.assertEqual([main.ay_farki(date(*a), BUGUN) for a in ((2026, 10, 12), (2026, 9, 12), (2026, 9, 13), (2026, 4, 12), (2026, 4, 13), (2025, 10, 12), (2024, 1, 15), (2026, 11, 1))], [0, 1, 0, 6, 5, 12, 32, 0])
        l = [d("6", date(2026, 4, 12)), d("7", date(2026, 3, 12)), d("12", date(2025, 10, 12)), d("13", date(2025, 9, 12)), d("24", date(2024, 10, 12)), d("25", date(2024, 9, 12)), d("yok", None), d("ileri", date(2026, 12, 1))]
        r = main.raporla(l, BUGUN)
        self.assertEqual({x["no"]: (x["yas"], x["dilim"]) for x in r["dosyalar"]},
                         {"6": (6, "0–6 ay"), "7": (7, "7–12 ay"), "12": (12, "7–12 ay"), "13": (13, "13–24 ay"), "24": (24, "13–24 ay"), "25": (25, "25+ ay"), "yok": (None, "(tarih yok)"), "ileri": (0, "0–6 ay")})
        self.assertEqual(([(x["dilim"], x["dosya"]) for x in r["dilimler"]], sorted((u["tur"], u["kim"]) for u in r["uyarilar"])),
                         ([("0–6 ay", 2), ("7–12 ay", 2), ("13–24 ay", 2), ("25+ ay", 1), ("(tarih yok)", 1)],
                          [("Tahsilatsız eski dosya", "12 · A"), ("Tahsilatsız eski dosya", "13 · A"), ("Tahsilatsız eski dosya", "24 · A"), ("Tahsilatsız eski dosya", "25 · A"), ("Takip tarihi ileride", "ileri · A")]))
        r = main.raporla(l[:6], BUGUN, dilimler=(12,), eski=25)
        self.assertEqual(([x["dilim"] for x in r["dilimler"]], [u["kim"] for u in r["uyarilar"]]), (["0–12 ay", "13+ ay"], ["25 · A"]))
        with self.assertRaises(ValueError):
            main.raporla([], BUGUN)

    def test_oranlar(self):
        l = [d("1", date(2024, 1, 15), 500, 350, 18, "X"), d("2", date(2026, 8, 20), 150, 150, 3, "x", "Kapalı"), d("3", date(2025, 2, 5), 800, 0, 12, "Y"), d("4", date(2025, 2, 5), 600, 0, 15, "Y", "Aciz"), d("5", date(2026, 4, 3), 100, 8, "9.5", "Z"),
             d("6", date(2026, 4, 3), 100, 120, 0, "Z")]
        r = main.raporla(l, BUGUN)
        self.assertEqual([(x["avukat"], x["dosya"], x["acik"], x["tutar"], x["tahsilat"], x["oran"], x["masraf"], x["net"], x["masraf_oran"], x["tahsilatsiz"], x["ort_yas"]) for x in r["avukatlar"]],
                         [("X", 2, 1, D(650), D(500), D("76.92"), D(21), D(479), D("4.20"), 0, D("16.50")), ("Z", 2, 2, D(200), D(128), D("64.00"), D("9.5"), D("118.5"), D("7.42"), 0, D("6.00")),
                          ("Y", 2, 1, D(1400), D(0), D("0.00"), D(27), D(-27), None, 2, D("20.00"))])
        self.assertEqual(r["toplam"], {"dosya": 6, "acik": 4, "tutar": D(2250), "tahsilat": D(628), "oran": D("27.91"), "masraf": D("57.5"), "net": D("570.5"), "masraf_oran": D("9.16"), "tahsilatsiz": 2, "ort_yas": D("14.17")})
        self.assertEqual([(m["avukat"], m["0–6 ay"]["dosya"], m["0–6 ay"]["oran"], m["13–24 ay"]["dosya"], m["13–24 ay"]["oran"], m["25+ ay"]["oran"]) for m in r["matris"]],
                         [("X", 1, D("100.00"), 0, None, D("70.00")), ("Z", 2, D("64.00"), 0, None, None), ("Y", 0, None, 2, D("0.00"), None)])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]), [("Bilgi", "Tahsilat takip tutarını aşıyor", "6 · Z"), ("Orta", "Masraf tahsilatı aşıyor", "5 · Z"), ("Orta", "Tahsilatsız eski dosya", "3 · Y")])
        x = {v["no"]: (v["oran"], v["net"]) for v in r["dosyalar"]}
        self.assertEqual((x["1"], x["5"], x["6"]), ((D("70.00"), D(332)), (D("8.00"), D("-1.5")), (D("120.00"), D(120))))


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "dosyalar.csv", cikti, BUGUN)
            self.assertEqual(r["toplam"], {"dosya": 9, "acik": 7, "tutar": D(3300000), "tahsilat": D(743000), "oran": D("22.52"), "masraf": D(80000), "net": D(663000), "masraf_oran": D("10.77"), "tahsilatsiz": 3, "ort_yas": D("13.33")})
            self.assertEqual([(x["avukat"], x["dosya"], x["acik"], x["tutar"], x["tahsilat"], x["oran"], x["masraf"], x["net"], x["masraf_oran"], x["tahsilatsiz"], x["ort_yas"]) for x in r["avukatlar"]],
                             [("Av. Kaya", 4, 3, D(1150000), D(590000), D("51.30"), D(32000), D(558000), D("5.42"), 1, D("14.00")), ("Av. Aksoy", 1, 1, D(100000), D(8000), D("8.00"), D(9500), D(-1500), D("118.75"), 0, D("6.00")),
                              ("Av. Demir", 4, 3, D(2050000), D(145000), D("7.07"), D(38500), D(106500), D("26.55"), 2, D("14.50"))])
            self.assertEqual([(x["dilim"], x["dosya"], x["tutar"], x["tahsilat"], x["oran"]) for x in r["dilimler"]],
                             [("0–6 ay", 3, D(500000), D(183000), D("36.60")), ("7–12 ay", 2, D(600000), D(120000), D("20.00")), ("13–24 ay", 3, D(1700000), D(90000), D("5.29")), ("25+ ay", 1, D(500000), D(350000), D("70.00"))])
            self.assertEqual({x["no"]: x["yas"] for x in r["dosyalar"]}, {"D-01": 32, "D-02": 16, "D-03": 7, "D-04": 1, "D-05": 20, "D-06": 11, "D-07": 3, "D-08": 24, "D-09": 6})
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]], [("Orta", "Masraf tahsilatı aşıyor", "D-09 · Av. Aksoy"), ("Orta", "Tahsilatsız eski dosya", "D-05 · Av. Demir")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Avukatlar", "Dosya Yaşı", "Avukat × Yaş", "Dosyalar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "d.csv"
            y.write_text("Esas No;Büro;Açılış Tarihi;Alacak Tutarı;Tahsil Edilen;Gider;Dosya Durumu\n2024/1;;15.01.2024;1.250.000,50;350.000;;İNFAZ\n2024/2;B;x;5;;;Aciz Vesikası\n2024/3;B;;0;;;\n2024/4;B;;5;-1;;\n2024/5;B;;5;;x;\n;B;;5;;;\n",
                         encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual((l, [(u["tur"], u["kim"]) for u in uy]),
                             ([{"no": "2024/1", "avukat": "(avukat yok)", "tarih": date(2024, 1, 15), "tutar": D("1250000.50"), "tahsilat": D(350000), "masraf": D(0), "durum": "Kapalı"},
                               {"no": "2024/2", "avukat": "B", "tarih": None, "tutar": D(5), "tahsilat": D(0), "masraf": D(0), "durum": "Aciz"}], [("Tarih okunamadı", "2024/2")] + [("Dosya okunamadı", f"2024/{n}") for n in (3, 4, 5)]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--dosyalar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "12.10.2026", "--dilimler", "12;24;36", "--eski", "18"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--dilimler", "24;12"], ["--eski", "0"])], [2] * 3)
            self.assertEqual(main.main(["--dosyalar", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Dosya No;Avukat\nA;B\n", "Dosya No;Takip Tutarı\n;1\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
