"""
Tahsilat ve Avukat Performans Raporu — Gerisi Bende kod bloğu
Krediler İzleme ve Takip › Yasal Takip Uzmanı

Takipteki alacakların tahsilat tutarlarını sözleşmeli avukat ve dosya yaşı bazında raporlar.
  Dosya yaşı (ay)      = takip tarihinden durum tarihine kadar geçen tam ay sayısı
  Yaş dilimi           = 0–6, 7–12, 13–24, 25+ ay   (--dilimler ile değişir)
  Tahsilat oranı %     = tahsilat ÷ takip tutarı × 100
  Net tahsilat         = tahsilat − masraf
  Masraf / tahsilat %  = masraf ÷ tahsilat × 100
  Tahsilatsız eski dosya = durumu açık, yaşı --eski (12) ay ve üstü, tahsilatı sıfır
Grup oranları toplamlar üzerinden hesaplanır (dosya oranlarının ortalaması değil). Avukatlar yalnız benzer yaş ve büyüklükteki dosyalarla karşılaştırılabilir.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
SUTUN = {"no": ("dosya no", "takip no", "esas no", "no"), "avukat": ("avukat", "buro", "hukuk burosu"), "tarih": ("takip tarihi", "takip baslangic tarihi", "acilis tarihi"), "tutar": ("takip tutari", "alacak tutari", "tutar"),
         "tahsilat": ("tahsilat", "tahsil edilen", "toplam tahsilat"), "masraf": ("masraf", "takip masrafi", "gider"), "durum": ("durum", "dosya durumu")}
KAPALI = {"kapali": "Kapalı", "kapandi": "Kapalı", "infaz": "Kapalı", "tahsil edildi": "Kapalı", "aciz": "Aciz", "aciz vesikasi": "Aciz", "derdest degil": "Kapalı"}
Z, YUZ = Decimal(0), Decimal(100)


def ay_farki(a: date, b: date) -> int:
    return max((b.year - a.year) * 12 + b.month - a.month - (1 if b.day < a.day else 0), 0)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("no", "tutar")):
        no = metin(r["no"])
        if not no:
            continue
        tutar, tahsilat, masraf = para(r["tutar"]), para(r["tahsilat"]) if metin(r["tahsilat"]) else Z, para(r["masraf"]) if metin(r["masraf"]) else Z
        if tutar is None or tutar <= 0 or tahsilat is None or tahsilat < 0 or masraf is None or masraf < 0:
            uy.append(uyari("Yüksek", "Dosya okunamadı", no, f"Satır {r['_satir']}: takip tutarı '{metin(r['tutar'])}', tahsilat '{metin(r['tahsilat'])}', masraf '{metin(r['masraf'])}'. Takip tutarı pozitif, diğerleri sıfır veya pozitif olmalı. "
                            "Satır alınmadı."))
            continue
        if metin(r["tarih"]) and tarih(r["tarih"]) is None:
            uy.append(uyari("Orta", "Tarih okunamadı", no, f"Takip tarihi '{metin(r['tarih'])}' okunamadı; dosya yaşı hesaplanamadı."))
        l.append({"no": no, "avukat": metin(r["avukat"]) or "(avukat yok)", "tarih": tarih(r["tarih"]), "tutar": tutar, "tahsilat": tahsilat, "masraf": masraf, "durum": KAPALI.get(katla(r["durum"]), "Açık")})
    return l, uy


def _oran(pay: Decimal, payda: Decimal) -> Decimal | None:
    return yuvarla(pay * YUZ / payda) if payda else None


def _topla(liste: list[dict]) -> dict:
    tutar, tahsilat, masraf = (sum((d[a] for d in liste), Z) for a in ("tutar", "tahsilat", "masraf"))
    yaslar = [d["yas"] for d in liste if d["yas"] is not None]
    return {"dosya": len(liste), "acik": sum(1 for d in liste if d["durum"] == "Açık"), "tutar": tutar, "tahsilat": tahsilat, "oran": _oran(tahsilat, tutar), "masraf": masraf, "net": tahsilat - masraf, "masraf_oran": _oran(masraf, tahsilat),
            "tahsilatsiz": sum(1 for d in liste if d["tahsilat"] == 0), "ort_yas": yuvarla(Decimal(sum(yaslar)) / len(yaslar)) if yaslar else None}


def raporla(dosyalar: list[dict], bugun: date, dilimler: tuple[int, ...] = (6, 12, 24), eski: int = 12) -> dict:
    if not dosyalar:
        raise ValueError("Raporlanacak dosya yok.")
    uy = []
    etiketler = [f"{a}–{b} ay" for a, b in zip((0,) + tuple(d + 1 for d in dilimler), dilimler)] + [f"{dilimler[-1] + 1}+ ay", "(tarih yok)"]
    for d in dosyalar:
        kim = f"{d['no']} · {d['avukat']}"
        d["yas"] = ay_farki(d["tarih"], bugun) if d["tarih"] else None
        d["dilim"] = etiketler[-1] if d["yas"] is None else etiketler[next((i for i, s in enumerate(dilimler) if d["yas"] <= s), len(dilimler))]
        d["oran"], d["net"] = _oran(d["tahsilat"], d["tutar"]), d["tahsilat"] - d["masraf"]
        if d["durum"] == "Açık" and d["yas"] is not None and d["yas"] >= eski and d["tahsilat"] == 0:
            uy.append(uyari("Orta", "Tahsilatsız eski dosya", kim, f"Takip {d['tarih']:%d.%m.%Y} tarihinde başladı ({d['yas']} ay); {d['tutar']:,.2f} tutarındaki dosyada tahsilat yok, masraf {d['masraf']:,.2f}. Dosyanın durumunu avukatla görüşün."))
        if d["masraf"] > d["tahsilat"] > 0:
            uy.append(uyari("Orta", "Masraf tahsilatı aşıyor", kim, f"Tahsilat {d['tahsilat']:,.2f}, masraf {d['masraf']:,.2f}; net {d['net']:,.2f}."))
        if d["tahsilat"] > d["tutar"]:
            uy.append(uyari("Bilgi", "Tahsilat takip tutarını aşıyor", kim, f"Takip tutarı {d['tutar']:,.2f}, tahsilat {d['tahsilat']:,.2f}. Faiz ve masraf tahsilatı dahil olabilir; takip tutarını aynı tanımla girin."))
        if d["tarih"] and d["tarih"] > bugun:
            uy.append(uyari("Orta", "Takip tarihi ileride", kim, f"Takip tarihi {d['tarih']:%d.%m.%Y}, durum tarihi {bugun:%d.%m.%Y}."))
    avukatlar = {}
    for d in dosyalar:
        avukatlar.setdefault(katla(d["avukat"]), []).append(d)
    av = sorted(({"avukat": g[0]["avukat"], **_topla(g)} for g in avukatlar.values()), key=lambda a: (-(a["oran"] or Z), katla(a["avukat"])))
    dilim = [{"dilim": e, **_topla([d for d in dosyalar if d["dilim"] == e])} for e in etiketler if e != "(tarih yok)" or any(d["dilim"] == e for d in dosyalar)]
    matris = [{"avukat": a["avukat"], **{e["dilim"]: _topla([d for d in avukatlar[katla(a["avukat"])] if d["dilim"] == e["dilim"]]) for e in dilim}} for a in av]
    return {"dosyalar": sorted(dosyalar, key=lambda d: (katla(d["avukat"]), -(d["yas"] or 0), katla(d["no"]))), "toplam": _topla(dosyalar), "avukatlar": av, "dilimler": dilim, "matris": matris, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, bugun: date) -> None:
    wb = Workbook()
    olcut = ["Dosya", "Açık Dosya", "Takip Tutarı", "Tahsilat", "Tahsilat Oranı %", "Masraf", "Net Tahsilat", "Masraf / Tahsilat %", "Tahsilatsız Dosya", "Ort. Dosya Yaşı (ay)"]
    deger = lambda o: [o["dosya"], o["acik"], o["tutar"], o["tahsilat"], o["oran"], o["masraf"], o["net"], o["masraf_oran"], o["tahsilatsiz"] or None, o["ort_yas"]]   # noqa: E731
    for i, (baslik, alan, ad) in enumerate((("Avukatlar", "avukatlar", "avukat"), ("Dosya Yaşı", "dilimler", "dilim"))):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = baslik
        _baslik(ws, ["Avukat" if i == 0 else "Dosya Yaşı"] + olcut, (26, 8, 9, 16, 16, 11, 14, 16, 11, 11, 11))
        for o in r[alan] + [{ad: "TOPLAM", **r["toplam"]}]:
            n = _satir(ws, [o[ad]] + deger(o), para_s=(4, 5, 6, 7, 8, 9, 11), dolgu={8: KIRMIZI if o["net"] < 0 else None})
        for h in ws[n]:
            h.font = Font(bold=True)
    ws.append([])
    ws.append([f"Durum tarihi: {bugun:%d.%m.%Y}"])
    ws = wb.create_sheet("Avukat × Yaş")
    adlar = [e["dilim"] for e in r["dilimler"]]
    _baslik(ws, ["Avukat"] + [f"{a} · dosya" for a in adlar] + [f"{a} · tahsilat %" for a in adlar], (26,) + (12,) * (2 * len(adlar)))
    for m in r["matris"]:
        _satir(ws, [m["avukat"]] + [m[a]["dosya"] or None for a in adlar] + [m[a]["oran"] for a in adlar], para_s=range(2 + len(adlar), 2 + 2 * len(adlar)))
    ws = wb.create_sheet("Dosyalar")
    _baslik(ws, ["Dosya No", "Avukat", "Takip Tarihi", "Yaş (ay)", "Yaş Dilimi", "Durum", "Takip Tutarı", "Tahsilat", "Tahsilat Oranı %", "Masraf", "Net Tahsilat"], (14, 24, 12, 9, 12, 10, 16, 16, 11, 14, 16))
    for d in r["dosyalar"]:
        _satir(ws, [d["no"], d["avukat"], d["tarih"], d["yas"], d["dilim"], d["durum"], d["tutar"], d["tahsilat"], d["oran"], d["masraf"], d["net"]], para_s=(7, 8, 9, 10, 11), tarih_s=(3,), dolgu={11: KIRMIZI if d["net"] < 0 else None})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, dilimler: tuple[int, ...] = (6, 12, 24), eski: int = 12) -> dict:
    dosyalar, uy = oku(yol)
    r = raporla(dosyalar, bugun, dilimler, eski)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, bugun)
    return r


def dilim_coz(s: str) -> tuple[int, ...] | None:
    try:
        d = tuple(int(x) for x in re.split(r"[;,\s]+", s.strip()) if x)
    except ValueError:
        return None
    return d if d and d[0] >= 0 and all(a < b for a, b in zip(d, d[1:])) else None


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Takipteki alacakların tahsilatını avukat ve dosya yaşı bazında raporlar.")
    p.add_argument("--dosyalar", type=Path, help="Dosya No, Avukat, Takip Tarihi, Takip Tutarı, Tahsilat, Masraf, Durum")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--dilimler", default="6;12;24", help="Dosya yaşı dilim sınırları, ay (varsayılan 6;12;24)")
    p.add_argument("--eski", type=int, default=12, help="Bu ay ve üstü yaşta, tahsilatı olmayan açık dosya uyarılır (varsayılan 12)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "tahsilat_avukat_performans.xlsx")
    a = p.parse_args(argv)
    ornek = a.dosyalar is None
    yol = ORNEK / "dosyalar.csv" if ornek else a.dosyalar
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    dilimler = dilim_coz(a.dilimler)
    if bugun is None or dilimler is None or a.eski < 1:
        print("[X] --tarih GG.AA.YYYY; --dilimler artan tam sayılar (örn. 6;12;24); --eski en az 1 olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, bugun, dilimler, a.eski)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {t['dosya']} dosya, {len(r['avukatlar'])} avukat · takip tutarı {t['tutar']:,.2f} · tahsilat {t['tahsilat']:,.2f} (%{t['oran']}) · masraf {t['masraf']:,.2f} · net {t['net']:,.2f}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
