"""
Zafiyet Tarama Önceliklendirme — Workers / Workless kod bloğu
Bilgi Teknolojileri › Bilgi Güvenliği Uzmanı

Zafiyet tarayıcısı çıktısını (Nessus, OpenVAS/Greenbone, Qualys vb. CSV/Excel) yalnız CVSS'e göre değil,
istismar bilgisi ve varlık bağlamıyla önceliklendirir:
  - CISA KEV (bilinen istismar edilen zafiyetler) kataloğu: KEV'deki CVE'ler en üst önceliğe çıkar; fidye yazılımı
    kampanyalarında kullanılanlar ayrıca işaretlenir.
  - FIRST EPSS skoru: önümüzdeki 30 günde istismar edilme olasılığı.
  - Varlık bağlamı: kritiklik (1–5) ve internete açıklık.
  - Öncelik (P1–P5) kural tablosuyla, düzeltme süresi (SLA) ayar dosyasıyla belirlenir; ilk görülme tarihi varsa
    süresi geçenler işaretlenir.
  - Aynı çözümle kapanan bulgular "düzeltme paketi" olarak gruplanır (hangi yama kaç sunucudaki kaç riski kapatır).
İnternete bağlanmaz: KEV ve EPSS dosyalarını kendiniz indirip verirsiniz.

Kullanım:
    python main.py                                          # kurgusal örnek verilerle
    python main.py --tarama nessus.csv --varliklar varliklar.xlsx --kev known_exploited_vulnerabilities.csv --epss epss_scores.csv
    python main.py --tarama openvas.csv --bugun 08.10.2026 --sla sla.json
"""
from __future__ import annotations

import argparse
import csv
import gzip
import json
import re
import sys
from collections import Counter, OrderedDict, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
CVE = re.compile(r"CVE-\d{4}-\d{4,7}", re.I)
VARSAYILAN_SLA = {"P1": 7, "P2": 15, "P3": 30, "P4": 90, "P5": 180}
SEVIYE_CVSS = {"critical": 9.5, "kritik": 9.5, "high": 8.0, "yüksek": 8.0, "medium": 5.5, "orta": 5.5, "low": 2.0, "düşük": 2.0,
               "info": 0.0, "none": 0.0, "bilgi": 0.0}


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def katla(s) -> str:
    s = kucuk(s)
    for a, b in zip("ıöüşğç", "iousgc"):
        s = s.replace(a, b)
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def sayi(x) -> float | None:
    if x in (None, ""):
        return None
    try:
        return float(str(x).replace(",", "."))
    except ValueError:
        return None


AYLAR_EN = {a: i for i, a in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


def tarih(x) -> date | None:
    """2026-10-08, 2026-10-08T10:00:00Z, 08.10.2026, 08/10/2026 veya Nessus biçimi 'Oct 8, 2026 10:00:00 UTC'."""
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = re.match(r"(\d{1,2})[./](\d{1,2})[./](\d{4})", s)
    if m:
        return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    m = re.match(r"([A-Za-z]{3})[a-z]* (\d{1,2}), (\d{4})", s)
    if m and m.group(1).lower() in AYLAR_EN:
        return date(int(m.group(3)), AYLAR_EN[m.group(1).lower()], int(m.group(2)))
    return None


def metin_oku(yol: Path) -> str:
    ham = gzip.decompress(yol.read_bytes()) if yol.suffix.lower() == ".gz" else yol.read_bytes()
    for kod in ("utf-8-sig", "cp1254", "latin-1"):
        try:
            return ham.decode(kod)
        except UnicodeDecodeError:
            continue
    return ""


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = metin_oku(yol)
        satirlar_m = [s for s in metin.splitlines() if not s.startswith("#")]      # EPSS dosyasındaki #model_version satırı
        ilk = "\n".join(satirlar_m[:10])
        satirlar = list(csv.reader(satirlar_m, delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def _bul(baslik, *adlar):
    b = [katla(x) for x in baslik]
    return next((b.index(katla(a)) for a in adlar if katla(a) in b), None)


def _al(r, i):
    return r[i] if i is not None and i < len(r) else None


# ----------------------------------------------------------------------------
# Okuma
# ----------------------------------------------------------------------------

def tarama_oku(yol: Path) -> list[dict]:
    """Nessus: Host, Port, Plugin ID, CVE, CVSS v3.0 Base Score, Risk, Name, Solution
    OpenVAS/Greenbone: IP, Hostname, Port, NVT OID, CVSS, Severity, NVT Name, CVEs, Solution
    Qualys: IP, DNS, QID, Title, Severity, CVE ID, CVSS3.1 Base, Solution"""
    s = tablo_oku(yol)
    b = s[0]
    i = {k: _bul(b, *v) for k, v in {
        "host": ("host", "ip", "ip address", "hostname", "dns", "varlık", "sunucu"),
        "ad_host": ("hostname", "dns", "netbios", "fqdn"),
        "port": ("port",), "id": ("plugin id", "nvt oid", "qid", "id", "bulgu id"),
        "cve": ("cve", "cves", "cve id", "cve ids"),
        "cvss": ("cvss v3 0 base score", "cvss v3 base score", "cvss3 1 base", "cvss3 base", "cvss v3", "cvss 3", "cvss base score", "cvss",
                 "cvss v2 0 base score"),
        "seviye": ("risk", "severity", "seviye", "risk factor"), "ad": ("name", "nvt name", "title", "zafiyet", "bulgu", "başlık"),
        "cozum": ("solution", "çözüm", "remediation", "fix"), "ilk": ("first seen", "first discovered", "ilk görülme", "first detected"),
    }.items()}
    if i["host"] is None or (i["cve"] is None and i["ad"] is None):
        raise SystemExit(f"Tarama dosyasında Host/IP ve CVE veya bulgu adı sütunları gerekli. Başlıklar: {b}")
    bulgular = []
    for r in s[1:]:
        host = str(_al(r, i["host"]) or "").strip()
        if not host:
            continue
        seviye = kucuk(_al(r, i["seviye"]))
        cvss = sayi(_al(r, i["cvss"]))
        if cvss is None:
            cvss = SEVIYE_CVSS.get(seviye)
        if (cvss or 0) == 0 and seviye in ("info", "none", "bilgi", "log"):
            continue                                            # bilgi amaçlı bulgular
        cveler = sorted({c.upper() for c in CVE.findall(str(_al(r, i["cve"]) or ""))})
        bulgular.append({"host": host, "ad_host": str(_al(r, i["ad_host"]) or "").strip(), "port": str(_al(r, i["port"]) or "").strip(),
                         "id": str(_al(r, i["id"]) or "").strip(), "cve": cveler, "cvss": cvss or 0.0,
                         "ad": str(_al(r, i["ad"]) or (cveler[0] if cveler else "")).strip(),
                         "cozum": str(_al(r, i["cozum"]) or "").strip(), "ilk": tarih(_al(r, i["ilk"]))})
    return bulgular


def varliklar_oku(yol: Path | None) -> dict[str, dict]:
    if not yol:
        return {}
    s = tablo_oku(yol)
    b = s[0]
    i_h = _bul(b, "host", "ip", "ip adresi", "hostname", "varlık", "sunucu")
    i_ad = _bul(b, "hostname", "ad", "varlık adı", "sunucu adı")
    i_k = _bul(b, "kritiklik", "kritiklik 1 5", "önem", "varlık değeri")
    i_i = _bul(b, "internete açık", "internet", "dmz", "dışa açık")
    i_s = _bul(b, "sahibi", "sorumlu", "ekip", "varlık sahibi")
    if i_h is None:
        raise SystemExit(f"Varlık dosyasında Host/IP sütunu gerekli. Başlıklar: {b}")
    sozel = {"cok yuksek": 5, "kritik": 5, "yuksek": 4, "orta": 3, "dusuk": 2, "cok dusuk": 1}
    v = {}
    for r in s[1:]:
        h = str(_al(r, i_h) or "").strip()
        if not h:
            continue
        k = sayi(_al(r, i_k))
        if k is None:
            k = sozel.get(katla(_al(r, i_k)), 3)
        kayit = {"kritiklik": int(max(1, min(5, k))), "internet": katla(_al(r, i_i)) in ("evet", "e", "1", "var", "yes", "true", "dmz"),
                 "sahip": str(_al(r, i_s) or "").strip(), "ad": str(_al(r, i_ad) or "").strip()}
        v[h] = kayit
        if kayit["ad"]:
            v.setdefault(kayit["ad"], kayit)
    return v


def kev_oku(yol: Path | None) -> dict[str, dict]:
    if not yol:
        return {}
    if yol.suffix.lower() == ".json":
        veri = json.loads(metin_oku(yol)).get("vulnerabilities", [])
        satirlar = [[x.get("cveID"), x.get("dateAdded"), x.get("dueDate"), x.get("knownRansomwareCampaignUse"),
                     x.get("vulnerabilityName")] for x in veri]
        b = ["cveID", "dateAdded", "dueDate", "knownRansomwareCampaignUse", "vulnerabilityName"]
    else:
        s = tablo_oku(yol)
        b, satirlar = s[0], s[1:]
    i_c, i_e, i_d = _bul(b, "cveid", "cve"), _bul(b, "dateadded", "eklenme"), _bul(b, "duedate", "son tarih")
    i_r, i_n = _bul(b, "knownransomwarecampaignuse", "ransomware"), _bul(b, "vulnerabilityname", "name")
    if i_c is None:
        raise SystemExit(f"KEV dosyasında cveID sütunu bulunamadı. Başlıklar: {b}")
    return {str(r[i_c]).strip().upper(): {"eklenme": tarih(_al(r, i_e)), "son": tarih(_al(r, i_d)),
                                          "fidye": katla(_al(r, i_r)) in ("known", "true", "evet", "1"), "ad": str(_al(r, i_n) or "")}
            for r in satirlar if _al(r, i_c)}


def epss_oku(yol: Path | None) -> dict[str, tuple[float, float]]:
    if not yol:
        return {}
    s = tablo_oku(yol)
    i_c, i_e, i_p = _bul(s[0], "cve"), _bul(s[0], "epss"), _bul(s[0], "percentile", "yüzdelik")
    if i_c is None or i_e is None:
        raise SystemExit(f"EPSS dosyasında cve ve epss sütunları gerekli. Başlıklar: {s[0]}")
    return {str(r[i_c]).strip().upper(): (sayi(_al(r, i_e)) or 0.0, sayi(_al(r, i_p)) or 0.0) for r in s[1:] if _al(r, i_c)}


# ----------------------------------------------------------------------------
# Önceliklendirme
# ----------------------------------------------------------------------------

def oncelik(cvss: float, kev: bool, fidye: bool, epss: float, kritiklik: int, internet: bool) -> tuple[str, str]:
    """Kural tablosu (sırayla ilk uyan): gerekçesiyle birlikte döner."""
    onemli = internet or kritiklik >= 4
    if kev and (onemli or fidye):
        return "P1", "KEV'de (aktif istismar)" + (" · fidye yazılımında kullanılıyor" if fidye else "") + \
               (" · internete açık" if internet else "") + (" · kritik varlık" if kritiklik >= 4 else "")
    if epss >= 0.5 and onemli and cvss >= 7:
        return "P1", f"EPSS %{epss * 100:.0f} · CVSS {cvss:g}" + (" · internete açık" if internet else " · kritik varlık")
    if kev:
        return "P2", "KEV'de (aktif istismar)"
    if cvss >= 9 and onemli:
        return "P2", f"CVSS {cvss:g} · " + ("internete açık" if internet else "kritik varlık")
    if epss >= 0.1 and cvss >= 7:
        return "P2", f"EPSS %{epss * 100:.0f} · CVSS {cvss:g}"
    if cvss >= 9 or (cvss >= 7 and onemli):
        return "P3", f"CVSS {cvss:g}" + (" · internete açık" if internet else " · kritik varlık" if kritiklik >= 4 else "")
    if cvss >= 7 or (cvss >= 4 and onemli):
        return "P4", f"CVSS {cvss:g}"
    return "P5", f"CVSS {cvss:g}"


def risk_puani(cvss: float, kev: bool, epss: float, kritiklik: int, internet: bool) -> float:
    """Sıralama için: CVSS × varlık ağırlığı × (1 + 4 × EPSS) × (KEV ise 2) × (internete açıksa 1,5)."""
    return round(cvss * (kritiklik / 3) * (1 + 4 * epss) * (2 if kev else 1) * (1.5 if internet else 1), 2)


def calistir(tarama: Path, cikti: Path, varlik_yolu: Path | None = None, kev_yolu: Path | None = None,
             epss_yolu: Path | None = None, sla: dict | None = None, bugun: date | None = None) -> dict:
    bulgular = tarama_oku(tarama)
    varliklar, kev, epss = varliklar_oku(varlik_yolu), kev_oku(kev_yolu), epss_oku(epss_yolu)
    sla = {**VARSAYILAN_SLA, **(sla or {})}
    bugun = bugun or date.today()
    uyarilar = []
    if not kev:
        uyarilar.append("KEV kataloğu verilmedi: aktif istismar bilgisi kullanılmadı (cisa.gov KEV CSV/JSON indirip --kev ile verin).")
    if not epss:
        uyarilar.append("EPSS dosyası verilmedi: istismar olasılığı kullanılmadı (first.org EPSS günlük CSV'si, --epss).")
    if not varliklar:
        uyarilar.append("Varlık listesi verilmedi: tüm varlıklar orta kritiklikte ve internete kapalı kabul edildi.")
    bilinmeyen = set()
    for f in bulgular:
        v = varliklar.get(f["host"]) or varliklar.get(f["ad_host"])
        if v is None:
            bilinmeyen.add(f["host"])
            v = {"kritiklik": 3, "internet": False, "sahip": "", "ad": ""}
        k = [kev[c] for c in f["cve"] if c in kev]
        e = max((epss.get(c, (0.0, 0.0)) for c in f["cve"]), default=(0.0, 0.0))
        f.update(kritiklik=v["kritiklik"], internet=v["internet"], sahip=v["sahip"], kev=bool(k), fidye=any(x["fidye"] for x in k),
                 kev_son=min((x["son"] for x in k if x["son"]), default=None), epss=e[0], epss_yuzdelik=e[1])
        f["oncelik"], f["gerekce"] = oncelik(f["cvss"], f["kev"], f["fidye"], f["epss"], f["kritiklik"], f["internet"])
        f["risk"] = risk_puani(f["cvss"], f["kev"], f["epss"], f["kritiklik"], f["internet"])
        bas = f["ilk"] or bugun
        f["hedef"] = bas + timedelta(days=sla[f["oncelik"]])
        f["gecikme"] = (bugun - f["hedef"]).days if f["ilk"] and bugun > f["hedef"] else 0
    if bilinmeyen and varliklar:
        uyarilar.append(f"Varlık listesinde olmayan {len(bilinmeyen)} host orta kritiklikte kabul edildi: {', '.join(sorted(bilinmeyen)[:10])}")
    bulgular.sort(key=lambda f: (f["oncelik"], -f["risk"], f["host"]))
    paketler = paketle(bulgular)
    _rapor(bulgular, paketler, sla, uyarilar, bugun, cikti)
    return {"bulgular": bulgular, "paketler": paketler, "uyarilar": uyarilar}


def paketle(bulgular: list[dict]) -> list[dict]:
    """Aynı bulgu (eklenti/QID veya ad) ve çözüm → tek düzeltme paketi."""
    gr: OrderedDict[tuple[str, str], dict] = OrderedDict()
    for f in bulgular:
        anahtar = (f["id"] or f["ad"], f["cozum"])
        p = gr.setdefault(anahtar, {"ad": f["ad"], "cozum": f["cozum"], "cve": set(), "hostlar": set(), "oncelik": f["oncelik"],
                                    "risk": 0.0, "kev": False, "internet": 0})
        p["cve"].update(f["cve"])
        p["hostlar"].add(f["host"])
        p["oncelik"] = min(p["oncelik"], f["oncelik"])
        p["risk"] += f["risk"]
        p["kev"] = p["kev"] or f["kev"]
        p["internet"] += 1 if f["internet"] else 0
    return sorted(gr.values(), key=lambda p: (p["oncelik"], -p["risk"]))


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
P_DOLGU = {"P1": "F8C9C6", "P2": "FDE2E1", "P3": "FFE8CC", "P4": "FFF4CE", "P5": "E8F0FE"}
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, satir=1):
    for c in ws[satir]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI


def _rapor(bulgular, paketler, sla, uyarilar, bugun, cikti):
    wb = Workbook()
    o = wb.active
    o.title = "Özet"
    o.append([f"Zafiyet önceliklendirme · {bugun:%d.%m.%Y}"])
    o["A1"].font = Font(bold=True, size=12)
    say = Counter(f["oncelik"] for f in bulgular)
    o.append(["Bulgu", len(bulgular)])
    o.append(["Varlık", len({f["host"] for f in bulgular})])
    o.append(["KEV'deki bulgu (aktif istismar)", sum(1 for f in bulgular if f["kev"])])
    o.append(["Fidye yazılımında kullanılan", sum(1 for f in bulgular if f["fidye"])])
    o.append(["SLA'sı geçmiş bulgu", sum(1 for f in bulgular if f["gecikme"] > 0)])
    o.append(["Düzeltme paketi", len(paketler)])
    o.append([])
    o.append(["Öncelik", "Bulgu", "Varlık", "Düzeltme Süresi (gün)", "Kural"])
    _baslik(o, o.max_row)
    kurallar = {"P1": "KEV + (internete açık veya kritik varlık veya fidye) · ya da EPSS ≥ %50, CVSS ≥ 7 ve internete açık/kritik",
                "P2": "KEV · ya da CVSS ≥ 9 ve internete açık/kritik · ya da EPSS ≥ %10 ve CVSS ≥ 7",
                "P3": "CVSS ≥ 9 · ya da CVSS ≥ 7 ve internete açık/kritik", "P4": "CVSS ≥ 7 · ya da CVSS ≥ 4 ve internete açık/kritik",
                "P5": "diğerleri"}
    for p in ("P1", "P2", "P3", "P4", "P5"):
        o.append([p, say.get(p, 0), len({f["host"] for f in bulgular if f["oncelik"] == p}), sla[p], kurallar[p]])
        o.cell(o.max_row, 1).fill = PatternFill("solid", fgColor=P_DOLGU[p])
    for u in uyarilar:
        o.append(["Uyarı", u])
    o.append([])
    o.append(["Not", "KEV'deki zafiyetlerde CISA'nın son tarihi (dueDate) ABD kamu kurumları içindir; burada yalnızca bilgi olarak "
                     "gösterilir. Öncelik kuralları ve SLA süreleri örnektir; kurumunuzun zafiyet yönetimi politikasına göre ayarlayın."])
    o.column_dimensions["A"].width = 32
    o.column_dimensions["B"].width = 10
    o.column_dimensions["C"].width = 10
    o.column_dimensions["D"].width = 20
    o.column_dimensions["E"].width = 100

    b = wb.create_sheet("Bulgular")
    b.append(["Öncelik", "Risk Puanı", "Host", "Port", "Bulgu", "CVE", "CVSS", "EPSS", "EPSS Yüzdelik", "KEV", "Fidye", "Kritiklik",
              "İnternete Açık", "Sahip", "Gerekçe", "İlk Görülme", "Hedef Kapanış", "Gecikme (gün)", "KEV Son Tarihi", "Çözüm", "Durum"])
    _baslik(b)
    for f in bulgular:
        b.append([f["oncelik"], f["risk"], f["host"], f["port"], f["ad"], ", ".join(f["cve"]), f["cvss"], f["epss"] or None,
                  f["epss_yuzdelik"] or None, "Evet" if f["kev"] else "", "Evet" if f["fidye"] else "", f["kritiklik"],
                  "Evet" if f["internet"] else "", f["sahip"], f["gerekce"], f["ilk"], f["hedef"], f["gecikme"] or None, f["kev_son"],
                  f["cozum"], None])
        n = b.max_row
        b.cell(n, 1).fill = PatternFill("solid", fgColor=P_DOLGU[f["oncelik"]])
        b.cell(n, 8).number_format = "0.0%"
        b.cell(n, 9).number_format = "0%"
        for c in (16, 17, 19):
            b.cell(n, c).number_format = "DD.MM.YYYY"
        if f["gecikme"]:
            b.cell(n, 18).fill = PatternFill("solid", fgColor=P_DOLGU["P1"])
    for j, w in enumerate((8, 9, 15, 7, 40, 22, 6, 7, 9, 6, 6, 9, 9, 14, 40, 11, 12, 9, 12, 50, 12), 1):
        b.column_dimensions[get_column_letter(j)].width = w
    b.freeze_panes = "C2"
    b.auto_filter.ref = b.dimensions

    p = wb.create_sheet("Düzeltme Planı")
    p.append(["Sıra", "Öncelik", "Bulgu", "Çözüm", "Etkilenen Host", "İnternete Açık Host", "CVE", "KEV", "Toplam Risk", "Sorumlu", "Planlanan Tarih"])
    _baslik(p)
    for i, x in enumerate(paketler, 1):
        p.append([i, x["oncelik"], x["ad"], x["cozum"], len(x["hostlar"]), x["internet"] or None, ", ".join(sorted(x["cve"]))[:200],
                  "Evet" if x["kev"] else "", round(x["risk"], 1), None, None])
        p.cell(p.max_row, 2).fill = PatternFill("solid", fgColor=P_DOLGU[x["oncelik"]])
        for c in (3, 4):
            p.cell(p.max_row, c).alignment = UST
    for j, w in enumerate((5, 8, 40, 55, 9, 10, 30, 6, 10, 14, 12), 1):
        p.column_dimensions[get_column_letter(j)].width = w
    p.freeze_panes = "C2"

    v = wb.create_sheet("Varlıklar")
    v.append(["Host", "Sahip", "Kritiklik", "İnternete Açık", "Bulgu", "P1", "P2", "KEV", "Toplam Risk"])
    _baslik(v)
    hostlar = defaultdict(list)
    for f in bulgular:
        hostlar[f["host"]].append(f)
    for h, fs in sorted(hostlar.items(), key=lambda x: -sum(f["risk"] for f in x[1])):
        v.append([h, fs[0]["sahip"], fs[0]["kritiklik"], "Evet" if fs[0]["internet"] else "", len(fs),
                  sum(1 for f in fs if f["oncelik"] == "P1") or None, sum(1 for f in fs if f["oncelik"] == "P2") or None,
                  sum(1 for f in fs if f["kev"]) or None, round(sum(f["risk"] for f in fs), 1)])
    for j, w in enumerate((16, 16, 9, 10, 7, 5, 5, 5, 11), 1):
        v.column_dimensions[get_column_letter(j)].width = w
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ornek = BURASI / "ornek_veri"
    ap = argparse.ArgumentParser(description="Zafiyet tarama çıktısını KEV, EPSS, CVSS ve varlık bağlamıyla önceliklendirir.")
    ap.add_argument("--tarama", type=Path, default=ornek / "tarama.csv", help="Tarayıcı çıktısı (Nessus/OpenVAS/Qualys CSV veya Excel)")
    ap.add_argument("--varliklar", type=Path, help="Varlık listesi: Host/IP, Kritiklik (1-5), İnternete Açık, Sahibi")
    ap.add_argument("--kev", type=Path, help="CISA KEV kataloğu (known_exploited_vulnerabilities.csv veya .json)")
    ap.add_argument("--epss", type=Path, help="FIRST EPSS skorları (epss_scores-AAAA-AA-GG.csv veya .csv.gz)")
    ap.add_argument("--sla", type=Path, help='Öncelik başına düzeltme süresi, gün (JSON: {"P1": 7, "P2": 15, ...})')
    ap.add_argument("--bugun", help="Değerlendirme tarihi (GG.AA.YYYY); varsayılan bugün")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "zafiyet_onceliklendirme.xlsx")
    a = ap.parse_args(argv)
    bugun = tarih(a.bugun) if a.bugun else None
    if a.tarama == ornek / "tarama.csv":
        a.varliklar = a.varliklar or ornek / "varliklar.csv"
        a.kev = a.kev or ornek / "kev_ornek.csv"
        a.epss = a.epss or ornek / "epss_ornek.csv"
        bugun = bugun or date(2026, 10, 8)
    sla = json.loads(a.sla.read_text(encoding="utf-8")) if a.sla else None
    s = calistir(a.tarama, a.cikti, a.varliklar, a.kev, a.epss, sla, bugun)
    say = Counter(f["oncelik"] for f in s["bulgular"])
    print(f"[OK] {len(s['bulgular'])} bulgu · " + " · ".join(f"{p}: {say.get(p, 0)}" for p in ("P1", "P2", "P3", "P4", "P5"))
          + f" · KEV: {sum(1 for f in s['bulgular'] if f['kev'])} · düzeltme paketi: {len(s['paketler'])}")
    for u in s["uyarilar"]:
        print(f"[!] {u}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
