"""
Ölçü Tablosu Beden Serisi (Grading) — Workers / Workless kod bloğu
Tekstil › Ürün Geliştirme / Modelhane

Ana bedenin (numune bedeni) ölçülerinden ve ölçü noktası bazındaki beden artış kurallarından tüm beden
serisinin ölçü tablosunu üretir. Artış her beden geçişi için ayrı verilebilir (ör. S→M 4 cm, L→XL 5 cm) ya da
tek değer olarak tüm geçişlere uygulanır. Ana bedenin üstüne artış eklenir, altına çıkarılır.
Toleransları tabloya ekler, artış kurallarındaki tutarsızlıkları (eksik geçiş, negatif ölçü, büyük bedenin
küçükten küçük olması) işaretler; istenirse inç karşılıklarını verir. İnternete bağlanmaz.

Kullanım:
    python main.py                                         # örnek erkek tişört kurallarıyla dener
    python main.py --girdi kurallar.xlsx --bedenler XS S M L XL XXL --ana M
    python main.py --girdi kurallar.xlsx --bedenler 36 38 40 42 44 46 --ana 40 --inc
"""
from __future__ import annotations

import argparse
import csv
import sys
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
INC = Decimal("2.54")


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def sayi(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace(" ", "").replace("±", "").replace("+/-", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def gecis_adi(a: str, b: str) -> str:
    return f"{a}-{b}"


def kurallari_oku(yol: Path, bedenler: list[str]) -> tuple[list[dict], list[str]]:
    s = tablo_oku(yol)
    b = [str(x or "").strip() for x in s[0]]
    bk = [kucuk(x) for x in b]
    i_ad = next((i for i, x in enumerate(bk) if x in ("ölçü noktası", "ölçü", "ölçü adı", "pom")), None)
    i_kod = next((i for i, x in enumerate(bk) if x in ("kod", "ölçü kodu", "pom kodu")), None)
    i_ana = next((i for i, x in enumerate(bk) if x in ("ana beden", "ana beden değeri", "numune", "numune bedeni", "değer")), None)
    i_tol = next((i for i, x in enumerate(bk) if x in ("tolerans", "tolerans (±)", "tol")), None)
    i_artis = next((i for i, x in enumerate(bk) if x in ("artış", "beden artışı", "grade", "fark")), None)
    if i_ad is None or i_ana is None:
        raise SystemExit(f"Kurallarda 'Ölçü Noktası' ve 'Ana Beden' sütunları gerekli. Başlıklar: {s[0]}")
    gecisler = [gecis_adi(bedenler[i], bedenler[i + 1]) for i in range(len(bedenler) - 1)]
    gecis_sutun = {g: next((i for i, x in enumerate(b) if x.replace(" ", "").replace("→", "-").replace(">", "-").upper() == g.upper()), None)
                   for g in gecisler}
    kurallar, uyarilar = [], []
    for r in s[1:]:
        ad = str(r[i_ad] or "").strip()
        if not ad:
            continue
        ana = sayi(r[i_ana])
        if ana is None:
            uyarilar.append(f"{ad}: ana beden değeri yok; atlandı")
            continue
        artislar = {}
        for g in gecisler:
            v = sayi(r[gecis_sutun[g]]) if gecis_sutun[g] is not None and gecis_sutun[g] < len(r) else None
            if v is None and i_artis is not None:
                v = sayi(r[i_artis])
            if v is None:
                uyarilar.append(f"{ad}: {g} geçişi için artış yok; 0 kabul edildi")
                v = Decimal(0)
            artislar[g] = v
        kurallar.append({"ad": ad, "kod": str(r[i_kod] or "").strip() if i_kod is not None else "", "ana": ana,
                         "tol": sayi(r[i_tol]) if i_tol is not None else None, "artis": artislar})
    return kurallar, uyarilar


def seri_uret(kural: dict, bedenler: list[str], ana: str) -> dict[str, Decimal]:
    i = bedenler.index(ana)
    deger = {ana: kural["ana"]}
    for j in range(i + 1, len(bedenler)):                      # yukarı: artış eklenir
        deger[bedenler[j]] = deger[bedenler[j - 1]] + kural["artis"][gecis_adi(bedenler[j - 1], bedenler[j])]
    for j in range(i - 1, -1, -1):                             # aşağı: artış çıkarılır
        deger[bedenler[j]] = deger[bedenler[j + 1]] - kural["artis"][gecis_adi(bedenler[j], bedenler[j + 1])]
    return {b: deger[b] for b in bedenler}


def kontrol(kural: dict, seri: dict[str, Decimal], bedenler: list[str]) -> list[str]:
    notlar = []
    if any(v <= 0 for v in seri.values()):
        notlar.append("sıfır/negatif ölçü")
    if any(seri[bedenler[i + 1]] < seri[bedenler[i]] for i in range(len(bedenler) - 1)):
        notlar.append("büyük beden küçükten küçük (negatif artış)")
    artislar = list(kural["artis"].values())
    if artislar and kural["tol"] is not None and any(a and abs(a) < kural["tol"] for a in artislar):
        notlar.append("artış toleranstan küçük: komşu bedenler ölçüyle ayırt edilemeyebilir")
    return notlar


def yuvarla(x: Decimal, adim: Decimal) -> Decimal:
    return (x / adim).quantize(Decimal(1), rounding=ROUND_HALF_UP) * adim


def calistir(girdi: Path, cikti: Path, bedenler: list[str], ana: str, inc: bool = False, adim: Decimal = Decimal("0.1")) -> dict:
    if ana not in bedenler:
        raise SystemExit(f"Ana beden ({ana}) beden serisinde yok: {bedenler}")
    if len(set(bedenler)) != len(bedenler):
        raise SystemExit("Beden serisinde tekrar eden beden var.")
    kurallar, uyarilar = kurallari_oku(girdi, bedenler)
    tablo = []
    for k in kurallar:
        seri = {b: yuvarla(v, adim) for b, v in seri_uret(k, bedenler, ana).items()}
        tablo.append({"kural": k, "seri": seri, "notlar": kontrol(k, seri, bedenler)})
    _rapor(tablo, bedenler, ana, cikti, inc, uyarilar)
    return {"tablo": tablo, "uyarilar": uyarilar}


BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
ANA_DOLGU = PatternFill("solid", fgColor="E8F0FE")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")


def _rapor(tablo, bedenler, ana, cikti, inc, uyarilar):
    wb = Workbook()
    ws = wb.active
    ws.title = "Ölçü Tablosu (cm)"
    ws.append(["Kod", "Ölçü Noktası", "Tolerans (±)"] + bedenler + ["Kontrol"])
    for h in ws[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    ai = 4 + bedenler.index(ana)
    ws.cell(1, ai).value = f"{ana} (ana)"
    for x in tablo:
        k = x["kural"]
        ws.append([k["kod"], k["ad"], None if k["tol"] is None else float(k["tol"])] + [float(x["seri"][b]) for b in bedenler]
                  + ["; ".join(x["notlar"]) or "Uygun"])
        ws.cell(ws.max_row, ai).fill = ANA_DOLGU
        if x["notlar"]:
            ws.cell(ws.max_row, 4 + len(bedenler)).fill = KIRMIZI
    for j, w in enumerate([8, 34, 12] + [10] * len(bedenler) + [50], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "D2"

    g = wb.create_sheet("Artış Kuralları")
    gecisler = [gecis_adi(bedenler[i], bedenler[i + 1]) for i in range(len(bedenler) - 1)]
    g.append(["Kod", "Ölçü Noktası", f"Ana ({ana})"] + gecisler)
    for h in g[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for x in tablo:
        k = x["kural"]
        g.append([k["kod"], k["ad"], float(k["ana"])] + [float(k["artis"][gc]) for gc in gecisler])
    for j, w in enumerate([8, 34, 12] + [10] * len(gecisler), 1):
        g.column_dimensions[get_column_letter(j)].width = w

    if inc:
        i = wb.create_sheet("Ölçü Tablosu (inç)")
        i.append(["Kod", "Ölçü Noktası", "Tolerans (±)"] + bedenler)
        for h in i[1]:
            h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        for x in tablo:
            k = x["kural"]
            i.append([k["kod"], k["ad"], None if k["tol"] is None else float((k["tol"] / INC).quantize(Decimal("0.01")))]
                     + [float((x["seri"][b] / INC).quantize(Decimal("0.01"))) for b in bedenler])
        for j, w in enumerate([8, 34, 12] + [10] * len(bedenler), 1):
            i.column_dimensions[get_column_letter(j)].width = w

    b = wb.create_sheet("Bilgi")
    bilgi = [["Yöntem", f"Ana beden ({ana}) değerine, üst bedenlere geçiş artışları eklenir, alt bedenlere çıkarılır"],
             ["Artış", "Her geçiş için ayrı sütun (ör. 'S-M') ya da tüm geçişlere uygulanan tek 'Artış' sütunu"],
             ["Kontrol", "Sıfır/negatif ölçü, büyük bedenin küçükten küçük olması, artışın toleranstan küçük olması"]]
    for s in bilgi + [["Uyarı", u] for u in uyarilar]:
        b.append(s)
    b.column_dimensions["A"].width = 12
    b.column_dimensions["B"].width = 120
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description="Ana beden ölçüleri ve artış kurallarından beden serisi ölçü tablosu üretir.")
    ap.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "erkek_tisort_kurallar.csv",
                    help="Kurallar (.xlsx/.csv): Ölçü Noktası, Ana Beden, [Kod, Tolerans, 'S-M' gibi geçiş sütunları veya tek 'Artış']")
    ap.add_argument("--bedenler", nargs="+", default=["XS", "S", "M", "L", "XL", "XXL"], help="Beden serisi, küçükten büyüğe")
    ap.add_argument("--ana", default="M", help="Ana (numune) beden")
    ap.add_argument("--inc", action="store_true", help="İnç karşılıklarını da ver")
    ap.add_argument("--adim", default="0.1", help="Yuvarlama adımı, cm (varsayılan 0.1; ör. 0.5)")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "olcu_tablosu.xlsx")
    a = ap.parse_args(argv)
    s = calistir(a.girdi, a.cikti, a.bedenler, a.ana, a.inc, sayi(a.adim) or Decimal("0.1"))
    sorunlu = [x for x in s["tablo"] if x["notlar"]]
    print(f"[OK] {len(s['tablo'])} ölçü noktası × {len(a.bedenler)} beden · kontrol uyarısı: {len(sorunlu)}")
    for x in sorunlu:
        print(f"[!] {x['kural']['ad']}: {'; '.join(x['notlar'])}")
    for u in s["uyarilar"]:
        print(f"[!] {u}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
