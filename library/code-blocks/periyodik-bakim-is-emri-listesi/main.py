"""
Periyodik Bakım İş Emri Listesi — Workers / Workless kod bloğu
Bakım Onarım › Bakım Teknisyeni

Bakım planından seçilen haftanın periyodik bakım iş emirlerini ve gereken malzemeleri çıkarır.
  Sonraki bakım tarihi = son bakım tarihi + periyot      (periyot: "30", "45 gün", "2 hafta", "3 ay", "1 yıl")
  Hafta                = --hafta ile verilen tarihin pazartesisinden pazarına
  İş emri açılır       = sonraki bakım tarihi hafta sonundan önce ya da hafta içindeyse
      Bu hafta         = tarih hafta içinde;   planlanan gün = sonraki bakım tarihi
      Gecikmiş         = tarih haftadan önce;  gecikme = hafta başı − sonraki bakım tarihi;   planlanan gün = hafta başı
      İlk bakım        = son bakım tarihi yazılmamış;   planlanan gün = hafta başı
  Malzeme ihtiyacı     = iş emirlerinin "Malzemeler" hücrelerindeki kalemlerin toplamı ("Gres x2, Filtre" biçiminde)
Periyotlar takvim süresidir; çalışma saati ya da sayaç bazlı bakımlar kapsam dışıdır.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
import calendar

SUTUN = {"ekipman": ("ekipman", "makine", "ekipman adi", "ekipman kodu"), "bakim": ("bakim adi", "bakim", "bakim tanimi", "is tanimi"), "periyot": ("periyot", "bakim periyodu", "siklik"),
         "son": ("son bakim tarihi", "son bakim", "son yapilma tarihi"), "sure": ("tahmini sure saat", "tahmini sure", "sure saat", "sure"), "sorumlu": ("sorumlu", "teknisyen", "sorumlu ekip"),
         "malzeme": ("malzemeler", "malzeme", "gereken malzeme")}
BIRIM = {"": ("gun", 1), "gun": ("gun", 1), "gunde": ("gun", 1), "hafta": ("gun", 7), "ay": ("ay", 1), "yil": ("ay", 12)}
DURUM_RENK = {"Gecikmiş": KIRMIZI, "İlk bakım": SARI, "Bu hafta": None}
SIFIR = Decimal(0)


def ay_ekle(t: date, n: int) -> date:
    y, m = divmod(t.year * 12 + t.month - 1 + n, 12)
    return date(y, m + 1, min(t.day, calendar.monthrange(y, m + 1)[1]))


def periyot_oku(x) -> tuple[str, int] | None:
    """'30' → ('gun', 30) · '2 hafta' → ('gun', 14) · '3 ay' → ('ay', 3) · '1 yıl' → ('ay', 12)"""
    m = re.fullmatch(r"(\d+)(?: 0)?\s*([a-z]*)", katla(x))
    if not m or m.group(2) not in BIRIM or int(m.group(1)) < 1:
        return None
    tur, kat = BIRIM[m.group(2)]
    return tur, int(m.group(1)) * kat


def sonraki(son: date, p: tuple[str, int]) -> date:
    return son + timedelta(days=p[1]) if p[0] == "gun" else ay_ekle(son, p[1])


def malzeme_oku(x) -> list[tuple[str, Decimal]]:
    l = []
    for parca in re.split(r"[,\n]", metin(x)):
        m = re.fullmatch(r"(.*?)(?:\s+[x×*]\s*(\d+(?:[.,]\d+)?))?", parca.strip(), re.I)
        if m and m.group(1).strip():
            l.append((m.group(1).strip(), Decimal(m.group(2).replace(",", ".")) if m.group(2) else Decimal(1)))
    return l


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("ekipman", "bakim", "periyot")):
        ekipman, bakim = metin(r["ekipman"]), metin(r["bakim"])
        if not ekipman:
            continue
        kim = f"{ekipman} · {bakim}"
        p, sure = periyot_oku(r["periyot"]), para(r["sure"]) if metin(r["sure"]) else SIFIR
        if p is None or sure is None or sure < 0 or not bakim:
            uy.append(uyari("Yüksek", "Bakım satırı okunamadı", kim, f"Satır {r['_satir']}: periyot '{metin(r['periyot'])}' (ör. 30, 2 hafta, 3 ay, 1 yıl), süre '{metin(r['sure'])}' ya da bakım adı okunamadı. Satır alınmadı."))
            continue
        if metin(r["son"]) and tarih(r["son"]) is None:
            uy.append(uyari("Orta", "Tarih okunamadı", kim, f"Son bakım tarihi '{metin(r['son'])}' okunamadı; hiç yapılmamış sayıldı."))
        l.append({"ekipman": ekipman, "bakim": bakim, "periyot": p, "periyot_yazi": metin(r["periyot"]), "son": tarih(r["son"]), "sure": sure, "sorumlu": metin(r["sorumlu"]) or "(sorumlu yok)", "malzeme": malzeme_oku(r["malzeme"])})
    return l, uy


def liste(plan: list[dict], gun: date, kritik: int = 30, kapasite: Decimal = Decimal(40)) -> dict:
    if not plan:
        raise ValueError("Bakım planında satır yok.")
    bas = gun - timedelta(days=gun.weekday())
    son, uy, emirler, gelecek = bas + timedelta(days=6), [], [], []
    for b in plan:
        kim = f"{b['ekipman']} · {b['bakim']}"
        b["sonraki"] = sonraki(b["son"], b["periyot"]) if b["son"] else None
        if b["sonraki"] is not None and b["sonraki"] > son:
            if b["sonraki"] <= son + timedelta(days=7):
                gelecek.append(b)
            continue
        e = dict(b)
        if b["sonraki"] is None:
            e.update(durum="İlk bakım", gecikme=None, gun=bas)
            uy.append(uyari("Orta", "Son bakım tarihi yok", kim, "Bakımın daha önce yapıldığına dair tarih yok; iş emri haftanın ilk gününe açıldı. Yapıldıktan sonra plana tarihini yazın."))
        elif b["sonraki"] < bas:
            e.update(durum="Gecikmiş", gecikme=(bas - b["sonraki"]).days, gun=bas)
            uy.append(uyari("Yüksek" if e["gecikme"] >= kritik else "Orta", "Gecikmiş bakım", kim, f"Bakım {b['sonraki']:%d.%m.%Y} tarihinde yapılmalıydı ({e['gecikme']} gün gecikme; son bakım {b['son']:%d.%m.%Y}, periyot {b['periyot_yazi']})."))
        else:
            e.update(durum="Bu hafta", gecikme=None, gun=b["sonraki"])
        emirler.append(e)
    emirler.sort(key=lambda e: (e["gun"], -(e["gecikme"] or 0), katla(e["ekipman"]), katla(e["bakim"])))
    iso = bas.isocalendar()
    for n, e in enumerate(emirler, 1):
        e["no"] = f"PB-{iso[0]}{iso[1]:02d}-{n:03d}"
    malzeme = {}
    for e in emirler:
        for ad, miktar in e["malzeme"]:
            m = malzeme.setdefault(katla(ad), {"ad": ad, "miktar": SIFIR, "emir": 0, "ekipman": []})
            m["miktar"], m["emir"] = m["miktar"] + miktar, m["emir"] + 1
            if e["ekipman"] not in m["ekipman"]:
                m["ekipman"].append(e["ekipman"])
    sorumlu = {}
    for e in emirler:
        s = sorumlu.setdefault(katla(e["sorumlu"]), {"ad": e["sorumlu"], "emir": 0, "sure": SIFIR, "gecikmis": 0})
        s["emir"], s["sure"], s["gecikmis"] = s["emir"] + 1, s["sure"] + e["sure"], s["gecikmis"] + (e["durum"] == "Gecikmiş")
    sorumlular = sorted(sorumlu.values(), key=lambda s: (-s["sure"], katla(s["ad"])))
    for s in sorumlular:
        if s["sure"] > kapasite:
            uy.append(uyari("Orta", "Haftalık yük kapasiteyi aşıyor", s["ad"], f"{s['emir']} iş emri, tahmini {s['sure']} saat (haftalık kapasite {kapasite} saat)."))
    gunler = [{"gun": bas + timedelta(days=i), "emir": sum(1 for e in emirler if e["gun"] == bas + timedelta(days=i)), "sure": sum((e["sure"] for e in emirler if e["gun"] == bas + timedelta(days=i)), SIFIR)} for i in range(7)]
    return {"bas": bas, "son": son, "emirler": emirler, "malzemeler": sorted(malzeme.values(), key=lambda m: katla(m["ad"])), "sorumlular": sorumlular, "gunler": gunler,
            "gelecek": sorted(gelecek, key=lambda b: (b["sonraki"], katla(b["ekipman"]))),
            "ozet": {"plan": len(plan), "emir": len(emirler), "gecikmis": sum(1 for e in emirler if e["durum"] == "Gecikmiş"), "ilk": sum(1 for e in emirler if e["durum"] == "İlk bakım"), "sure": sum((e["sure"] for e in emirler), SIFIR),
                     "gelecek": len(gelecek)}, "uyarilar": uy}


GUN_AD = ("Pazartesi", "Salı", "Çarşamba", "Perşembe", "Cuma", "Cumartesi", "Pazar")


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "İş Emirleri"
    _baslik(ws, ["İş Emri No", "Durum", "Planlanan Gün", "Ekipman", "Bakım", "Periyot", "Son Bakım", "Olması Gereken Tarih", "Gecikme (gün)", "Tahmini Süre (saat)", "Sorumlu", "Malzemeler", "Yapıldı (tarih)", "Not"],
            (16, 11, 12, 24, 32, 10, 12, 12, 9, 10, 20, 36, 12, 28))
    for e in r["emirler"]:
        _satir(ws, [e["no"], e["durum"], e["gun"], e["ekipman"], e["bakim"], e["periyot_yazi"], e["son"], e["sonraki"], e["gecikme"], e["sure"] or None, e["sorumlu"],
                    ", ".join(f"{ad} x{m.normalize():f}" for ad, m in e["malzeme"]), None, None], tarih_s=(3, 7, 8), dolgu={2: DURUM_RENK[e["durum"]]}, sar=(5, 12))
    ws.append([])
    ws.append([f"Hafta: {r['bas']:%d.%m.%Y} – {r['son']:%d.%m.%Y}"])
    ws = wb.create_sheet("Malzeme İhtiyacı")
    _baslik(ws, ["Malzeme", "Toplam Miktar", "İş Emri", "Ekipmanlar"], (34, 13, 9, 60))
    for m in r["malzemeler"]:
        _satir(ws, [m["ad"], m["miktar"], m["emir"], ", ".join(m["ekipman"])], sar=(4,))
    ws = wb.create_sheet("Haftalık Yük")
    _baslik(ws, ["Gün", "Tarih", "İş Emri", "Tahmini Süre (saat)"], (14, 12, 9, 12))
    for i, g in enumerate(r["gunler"]):
        _satir(ws, [GUN_AD[i], g["gun"], g["emir"] or None, g["sure"] or None], tarih_s=(2,))
    ws.append([])
    _baslik(ws, ["Sorumlu", "İş Emri", "Gecikmiş", "Tahmini Süre (saat)"], dondur="A2")
    for s in r["sorumlular"]:
        _satir(ws, [s["ad"], s["emir"], s["gecikmis"] or None, s["sure"]])
    ws = wb.create_sheet("Sonraki Hafta")
    _baslik(ws, ["Olması Gereken Tarih", "Ekipman", "Bakım", "Periyot", "Son Bakım", "Sorumlu"], (14, 24, 32, 10, 12, 20))
    for b in r["gelecek"]:
        _satir(ws, [b["sonraki"], b["ekipman"], b["bakim"], b["periyot_yazi"], b["son"], b["sorumlu"]], tarih_s=(1, 5))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, gun: date, kritik: int = 30, kapasite: Decimal = Decimal(40)) -> dict:
    plan, uy = oku(yol)
    r = liste(plan, gun, kritik, kapasite)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Bakım planından haftalık periyodik bakım iş emirlerini ve gereken malzemeleri çıkarır.")
    p.add_argument("--plan", type=Path, help="Ekipman, Bakım Adı, Periyot, Son Bakım Tarihi, Tahmini Süre (saat), Sorumlu, Malzemeler")
    p.add_argument("--hafta", help="İş emirleri çıkarılacak haftadan bir gün (GG.AA.YYYY); verilmezse bugünün haftası")
    p.add_argument("--kritik", type=int, default=30, help="Bu kadar gün ve üstü gecikme yüksek önemli sayılır (varsayılan 30)")
    p.add_argument("--kapasite", default="40", help="Sorumlu başına haftalık saat kapasitesi; aşılırsa uyarılır (varsayılan 40)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "periyodik_bakim_is_emirleri.xlsx")
    a = p.parse_args(argv)
    ornek = a.plan is None
    yol = ORNEK / "bakim_plani.csv" if ornek else a.plan
    gun, kapasite = tarih(a.hafta) if a.hafta else date(2026, 10, 12) if ornek else date.today(), para(a.kapasite)
    if gun is None or a.kritik < 1 or kapasite is None or kapasite <= 0:
        print("[X] --hafta GG.AA.YYYY; --kritik en az 1; --kapasite sıfırdan büyük olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, gun, a.kritik, kapasite)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] Hafta {r['bas']:%d.%m.%Y}–{r['son']:%d.%m.%Y} · {o['emir']} iş emri (gecikmiş {o['gecikmis']}, ilk bakım {o['ilk']}) · tahmini {o['sure']} saat · {len(r['malzemeler'])} malzeme kalemi · sonraki hafta {o['gelecek']} bakım")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
