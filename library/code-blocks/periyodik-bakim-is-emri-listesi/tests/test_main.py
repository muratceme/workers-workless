import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
PZT = date(2026, 10, 12)


def b(ekipman, bakim, periyot, son, sure=1, sorumlu="S", malzeme=""):
    return {"ekipman": ekipman, "bakim": bakim, "periyot": main.periyot_oku(periyot), "periyot_yazi": periyot, "son": son, "sure": D(str(sure)), "sorumlu": sorumlu, "malzeme": main.malzeme_oku(malzeme)}


def veri():
    return [b("E1", "Yağlama", "30", date(2026, 9, 14), 1, "A", "Gres x2, Bez"), b("E1", "Filtre", "3 ay", date(2026, 7, 10), 2, "A", "Filtre x1, gres x1"), b("E2", "Genel", "1 yıl", date(2025, 10, 18), 4, "B"),
            b("E2", "Kayış", "2 hafta", date(2026, 10, 5)), b("E3", "Kalibrasyon", "6 ay", None, 1.5, "B"), b("E3", "Yağ", "90 gün", date(2026, 6, 1), 0.5, "A"), b("E4", "X", "30", date(2026, 10, 1))]


class ListeTesti(unittest.TestCase):
    def test_is_emirleri(self):
        for gun in (PZT, date(2026, 10, 15), date(2026, 10, 18)):
            r = main.liste(veri(), gun)
            self.assertEqual((r["bas"], r["son"]), (PZT, date(2026, 10, 18)))
        self.assertEqual([(e["no"], e["ekipman"], e["bakim"], e["durum"], e["gun"], e["sonraki"], e["gecikme"]) for e in r["emirler"]],
                         [("PB-202642-001", "E3", "Yağ", "Gecikmiş", PZT, date(2026, 8, 30), 43), ("PB-202642-002", "E1", "Filtre", "Gecikmiş", PZT, date(2026, 10, 10), 2), ("PB-202642-003", "E3", "Kalibrasyon", "İlk bakım", PZT, None, None),
                          ("PB-202642-004", "E1", "Yağlama", "Bu hafta", date(2026, 10, 14), date(2026, 10, 14), None), ("PB-202642-005", "E2", "Genel", "Bu hafta", date(2026, 10, 18), date(2026, 10, 18), None)])
        self.assertEqual([(x["ekipman"], x["bakim"], x["sonraki"]) for x in r["gelecek"]], [("E2", "Kayış", date(2026, 10, 19))])
        self.assertEqual([(m["ad"], m["miktar"], m["emir"], m["ekipman"]) for m in r["malzemeler"]], [("Bez", D(1), 1, ["E1"]), ("Filtre", D(1), 1, ["E1"]), ("gres", D(3), 2, ["E1"])])
        self.assertEqual(r["sorumlular"], [{"ad": "B", "emir": 2, "sure": D("5.5"), "gecikmis": 0}, {"ad": "A", "emir": 3, "sure": D("3.5"), "gecikmis": 2}])
        self.assertEqual([(g["gun"].day, g["emir"], g["sure"]) for g in r["gunler"]], [(12, 3, D("4.0")), (13, 0, D(0)), (14, 1, D(1)), (15, 0, D(0)), (16, 0, D(0)), (17, 0, D(0)), (18, 1, D(4))])
        self.assertEqual((r["ozet"], sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"])),
                         ({"plan": 7, "emir": 5, "gecikmis": 2, "ilk": 1, "sure": D("9.0"), "gelecek": 1}, [("Orta", "Gecikmiş bakım", "E1 · Filtre"), ("Orta", "Son bakım tarihi yok", "E3 · Kalibrasyon"), ("Yüksek", "Gecikmiş bakım", "E3 · Yağ")]))

    def test_secenekler_ve_ayristirma(self):
        u = lambda **kw: sorted((x["onem"], x["tur"]) for x in main.liste(veri(), PZT, **kw)["uyarilar"] if x["tur"] != "Son bakım tarihi yok")   # noqa: E731
        self.assertEqual((u(kritik=44), u(kritik=2, kapasite=D(5))), ([("Orta", "Gecikmiş bakım")] * 2, [("Orta", "Haftalık yük kapasiteyi aşıyor"), ("Yüksek", "Gecikmiş bakım"), ("Yüksek", "Gecikmiş bakım")]))
        self.assertEqual([main.periyot_oku(x) for x in ("30", "45 gün", "2 hafta", "3 Ay", "1 YIL", 7, 30.0, "0", "ay", "3 saat", "", "x")], [("gun", 30), ("gun", 45), ("gun", 14), ("ay", 3), ("ay", 12), ("gun", 7), ("gun", 30), None, None, None, None, None])
        self.assertEqual((main.sonraki(date(2026, 1, 31), ("ay", 1)), main.sonraki(date(2024, 2, 29), ("ay", 12)), main.sonraki(date(2026, 12, 20), ("gun", 14))), (date(2026, 2, 28), date(2025, 2, 28), date(2027, 1, 3)))
        self.assertEqual(main.malzeme_oku("Gres (kg) x2, Bez , Filtre 10x20 x 3, Yağ × 1,5\nRulman 6204, , Conta*4"),
                         [("Gres (kg)", D(2)), ("Bez", D(1)), ("Filtre 10x20", D(3)), ("Yağ", D(1)), ("5", D(1)), ("Rulman 6204", D(1)), ("Conta*4", D(1))])
        self.assertEqual(main.liste([b("E", "B", "30", date(2026, 10, 1))], PZT)["emirler"], [])
        with self.assertRaises(ValueError):
            main.liste([], PZT)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "bakim_plani.csv", cikti, PZT)
            self.assertEqual([(e["no"][-3:], e["ekipman"], e["durum"], e["gun"].day, e["gecikme"]) for e in r["emirler"]],
                             [("001", "CNC 1", "Gecikmiş", 12, 53), ("002", "Kompresör", "İlk bakım", 12, None), ("003", "Pres 2", "Bu hafta", 12, None), ("004", "Pres 1", "Bu hafta", 13, None), ("005", "Pres 1", "Bu hafta", 14, None),
                              ("006", "Ana pano", "Bu hafta", 15, None), ("007", "Forklift", "Bu hafta", 15, None), ("008", "Pres 2", "Bu hafta", 16, None)])
            self.assertEqual([(m["ad"], m["miktar"], m["emir"], m["ekipman"]) for m in r["malzemeler"]],
                             [("Gres (kg)", D(2), 1, ["Pres 2"]), ("Hava filtresi", D(1), 1, ["Kompresör"]), ("Hidrolik filtre", D(1), 1, ["Pres 1"]), ("Hidrolik yağ (lt)", D(20), 1, ["Pres 1"]), ("Saf su (lt)", D(2), 1, ["Forklift"]),
                              ("Soğutma sıvısı (lt)", D(60), 1, ["CNC 1"]), ("Temizlik bezi", D(3), 2, ["CNC 1", "Pres 2"])])
            self.assertEqual([(s["ad"], s["emir"], s["sure"], s["gecikmis"]) for s in r["sorumlular"]], [("Mekanik Bakım", 5, D("8.0"), 1), ("Elektrik Bakım", 3, D("3.5"), 0)])
            self.assertEqual([(g["emir"], g["sure"]) for g in r["gunler"]], [(3, D("5.5")), (1, D("0.5")), (1, D(2)), (2, D("2.5")), (1, D(1)), (0, D(0)), (0, D(0))])
            self.assertEqual(([(x["ekipman"], x["sonraki"].day) for x in r["gelecek"]], r["ozet"]["sure"], [(u["onem"], u["tur"], u["kim"][:9]) for u in r["uyarilar"]]),
                             ([("CNC 1", 19), ("Kompresör", 20)], D("11.5"), [("Yüksek", "Gecikmiş bakım", "CNC 1 · S"), ("Orta", "Son bakım tarihi yok", "Kompresör")]))
            self.assertEqual(load_workbook(cikti).sheetnames, ["İş Emirleri", "Malzeme İhtiyacı", "Haftalık Yük", "Sonraki Hafta", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "p.csv"
            y.write_text("Makine;Bakım;Sıklık;Son Bakım;Süre;Teknisyen;Malzeme\nM1;Yağlama;2 hafta;x;1,5;Ali;Gres x2\nM1;Kontrol;7;01.10.2026;;;\nM2;Y;3 saat;;;;\nM2;;30;;;;\nM3;Z;30;;-1;;\n;Z;30;;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["ekipman"], x["bakim"], x["periyot"], x["son"], x["sure"], x["sorumlu"], x["malzeme"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("M1", "Yağlama", ("gun", 14), None, D("1.5"), "Ali", [("Gres", D(2))]), ("M1", "Kontrol", ("gun", 7), date(2026, 10, 1), D(0), "(sorumlu yok)", [])],
                              [("Tarih okunamadı", "M1 · Yağlama"), ("Bakım satırı okunamadı", "M2 · Y"), ("Bakım satırı okunamadı", "M2 ·"), ("Bakım satırı okunamadı", "M3 · Z")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--plan", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--hafta", "19.10.2026", "--kritik", "14", "--kapasite", "30"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--hafta", "x"], ["--kritik", "0"], ["--kapasite", "0"])], [2] * 3)
            self.assertEqual(main.main(["--plan", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Ekipman;Bakım Adı\nA;B\n", "Ekipman;Bakım Adı;Periyot\n;B;30\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
