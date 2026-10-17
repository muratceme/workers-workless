"""
Yapılandırma Takip Listesi — Gerisi Bende kod bloğu
Krediler İzleme ve Takip › Kredi İzleme Uzmanı

Yeniden yapılandırılan kredilerin ödeme planı uyumunu ve izleme süresini izler.
  Ödeme uyumu %     = ödenen tutar ÷ vadesi gelen tutar × 100   (vadesi gelen taksit yoksa hesaplanmaz)
  Eksik ödeme       = vadesi gelen tutar − ödenen tutar   (eksiye düşmez)
  İzleme bitişi     = yapılandırma tarihi + izleme süresi (ay)   (satırda yoksa --izleme-ay, varsayılan 12)
  Durum             = Uyumsuz             : gecikme günü --gecikme (30) günü aşmış
                      Gecikmeli           : gecikme var ya da eksik ödeme var
                      İzleme süresi doldu : gecikme ve eksik ödeme yok, izleme bitişi geçmiş
                      Uyumlu              : diğerleri
İzleme süresi ve gecikme eşiği çalışma varsayımıdır; sınıflandırma kararı ilgili düzenlemelere ve izleme biriminizin değerlendirmesine göre verilir.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
import calendar

SUTUN = {"no": ("musteri no", "musteri numarasi"), "ad": ("musteri", "musteri adi", "unvan"), "kredi": ("kredi no", "hesap no", "kredi numarasi"), "ytarih": ("yapilandirma tarihi", "yapilandirma"),
         "ytutar": ("yapilandirilan tutar", "yapilandirma tutari"), "risk": ("kalan risk", "risk bakiyesi", "risk", "bakiye"), "vadesi": ("vadesi gelen tutar", "vadesi gelen", "tahakkuk eden"), "odenen": ("odenen tutar", "odenen", "tahsil edilen"),
         "gun": ("gecikme gun", "gecikme gunu", "gecikme"), "izleme": ("izleme suresi ay", "izleme suresi", "izleme ay"), "sayi": ("yapilandirma sayisi", "kacinci yapilandirma")}
DURUMLAR = ("Uyumsuz", "Gecikmeli", "İzleme süresi doldu", "Uyumlu")
DURUM_RENK = {"Uyumsuz": KIRMIZI, "Gecikmeli": SARI, "İzleme süresi doldu": MAVI, "Uyumlu": YESIL}
Z, YUZ = Decimal(0), Decimal(100)


def ay_ekle(t: date, n: int) -> date:
    y, a = divmod(t.year * 12 + t.month - 1 + n, 12)
    return date(y, a + 1, min(t.day, calendar.monthrange(y, a + 1)[1]))


def _tam(x, varsayilan=None) -> int | None:
    if not metin(x):
        return varsayilan
    d = para(x)
    return int(d) if d is not None and d >= 0 and d == d.to_integral_value() else -1


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("no", "ytarih")):
        no = metin(r["no"])
        if not no:
            continue
        kim = f"{no} · {metin(r['ad'])} · {metin(r['kredi'])}".replace(" ·  ·", " ·").strip(" ·")
        d = {a: (para(r[a]) if metin(r[a]) else Z) for a in ("ytutar", "risk", "vadesi", "odenen")}
        gun, izleme, sayi, yt = _tam(r["gun"], 0), _tam(r["izleme"]), _tam(r["sayi"], 1), tarih(r["ytarih"])
        if yt is None or any(v is None or v < 0 for v in d.values()) or -1 in (gun, izleme, sayi) or izleme == 0:
            uy.append(uyari("Yüksek", "Satır okunamadı", kim, f"Satır {r['_satir']}: yapılandırma tarihi, tutarlar, gecikme günü, izleme süresi ya da yapılandırma sayısı okunamadı (eksi ya da sayı olmayan değer). Satır alınmadı."))
            continue
        l.append({"no": no, "ad": metin(r["ad"]), "kredi": metin(r["kredi"]), "ytarih": yt, **d, "gun": gun, "izleme": izleme, "sayi": max(sayi, 1)})
    return l, uy


def degerlendir(krediler: list[dict], bugun: date, gecikme: int = 30, izleme_ay: int = 12, gun_once: int = 30) -> dict:
    if not krediler:
        raise ValueError("İzlenecek yapılandırılmış kredi yok.")
    uy = []
    for k in krediler:
        kim = f"{k['no']} · {k['ad']} · {k['kredi']}".replace(" ·  ·", " ·").strip(" ·")
        k["izleme_ay"] = k["izleme"] or izleme_ay
        k["izleme_bitis"] = ay_ekle(k["ytarih"], k["izleme_ay"])
        k["izleme_kalan"] = (k["izleme_bitis"] - bugun).days
        k["uyum"] = yuvarla(k["odenen"] * YUZ / k["vadesi"]) if k["vadesi"] else None
        k["eksik"] = max(k["vadesi"] - k["odenen"], Z)
        if k["gun"] > gecikme:
            k["durum"] = "Uyumsuz"
            uy.append(uyari("Yüksek", "Ödeme planına uyumsuz", kim, f"Gecikme {k['gun']} gün (eşik {gecikme}). Vadesi gelen {k['vadesi']:,.2f}, ödenen {k['odenen']:,.2f}; eksik {k['eksik']:,.2f}. Kalan risk {k['risk']:,.2f}. "
                            "Sınıflandırmayı izleme biriminizle değerlendirin."))
        elif k["gun"] > 0 or k["eksik"] > 0:
            k["durum"] = "Gecikmeli"
            uy.append(uyari("Orta", "Gecikmeli ödeme", kim, f"Gecikme {k['gun']} gün; eksik ödeme {k['eksik']:,.2f}. Eşiğe ({gecikme} gün) {gecikme - k['gun']} gün kaldı."))
        elif k["izleme_kalan"] <= 0:
            k["durum"] = "İzleme süresi doldu"
            uy.append(uyari("Bilgi", "İzleme süresi doldu", kim, f"{k['izleme_ay']} aylık izleme süresi {k['izleme_bitis']:%d.%m.%Y} tarihinde doldu; gecikme ve eksik ödeme yok. Sınıf değişikliği koşulları izleme biriminizce değerlendirilebilir."))
        else:
            k["durum"] = "Uyumlu"
            if k["izleme_kalan"] <= gun_once:
                uy.append(uyari("Bilgi", "İzleme süresi bitiyor", kim, f"İzleme süresi {k['izleme_bitis']:%d.%m.%Y} tarihinde doluyor ({k['izleme_kalan']} gün kaldı)."))
        if k["sayi"] > 1:
            uy.append(uyari("Orta", "Birden çok kez yapılandırılmış", kim, f"Bu kredi {k['sayi']}. kez yapılandırılmış. Kalan risk {k['risk']:,.2f}."))
        if (k["gun"] > 0) != (k["eksik"] > 0):
            uy.append(uyari("Orta", "Gecikme bilgisi tutarsız", kim, f"Gecikme günü {k['gun']}, eksik ödeme {k['eksik']:,.2f}; biri sıfır, diğeri değil. Veriyi kontrol edin."))
        if k["ytarih"] > bugun:
            uy.append(uyari("Orta", "Yapılandırma tarihi ileride", kim, f"Yapılandırma tarihi {k['ytarih']:%d.%m.%Y}, durum tarihi {bugun:%d.%m.%Y}."))
    krediler.sort(key=lambda k: (DURUMLAR.index(k["durum"]), -k["risk"], katla(k["no"])))
    toplam = sum((k["risk"] for k in krediler), Z)
    ozet = [{"durum": d, "kredi": sum(1 for k in krediler if k["durum"] == d), "risk": sum((k["risk"] for k in krediler if k["durum"] == d), Z)} for d in DURUMLAR]
    for o in ozet:
        o["pay"] = yuvarla(o["risk"] * YUZ / toplam) if toplam else None
    vadesi, odenen = sum((k["vadesi"] for k in krediler), Z), sum((k["odenen"] for k in krediler), Z)
    return {"krediler": krediler, "ozet": ozet, "toplam": {"risk": toplam, "ytutar": sum((k["ytutar"] for k in krediler), Z), "vadesi": vadesi, "odenen": odenen, "eksik": sum((k["eksik"] for k in krediler), Z),
                                                         "uyum": yuvarla(odenen * YUZ / vadesi) if vadesi else None}, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, bugun: date) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Yapılandırmalar"
    _baslik(ws, ["Durum", "Müşteri No", "Müşteri", "Kredi No", "Yapılandırma Tarihi", "Yapılandırılan Tutar", "Kalan Risk", "Vadesi Gelen", "Ödenen", "Eksik Ödeme", "Uyum %", "Gecikme (gün)", "İzleme Bitişi", "İzlemeye Kalan Gün",
                 "Yapılandırma Sayısı", "Not (siz doldurun)"], (20, 12, 26, 12, 13, 16, 16, 15, 15, 15, 9, 10, 12, 11, 11, 30))
    for k in r["krediler"]:
        _satir(ws, [k["durum"], k["no"], k["ad"], k["kredi"], k["ytarih"], k["ytutar"], k["risk"], k["vadesi"], k["odenen"], k["eksik"] or None, k["uyum"], k["gun"] or None, k["izleme_bitis"], k["izleme_kalan"], k["sayi"], None],
               para_s=(6, 7, 8, 9, 10, 11), tarih_s=(5, 13), dolgu={1: DURUM_RENK[k["durum"]], 15: SARI if k["sayi"] > 1 else None})
    ws = wb.create_sheet("Özet")
    _baslik(ws, ["Durum", "Kredi", "Kalan Risk", "Pay %"], (24, 8, 18, 9))
    for o in r["ozet"]:
        _satir(ws, [o["durum"], o["kredi"], o["risk"], o["pay"]], para_s=(3, 4), dolgu={1: DURUM_RENK[o["durum"]]})
    t = r["toplam"]
    n = _satir(ws, ["TOPLAM", len(r["krediler"]), t["risk"], None], para_s=(3,))
    for h in ws[n]:
        h.font = Font(bold=True)
    ws.append([])
    for etiket, deger in (("Yapılandırılan tutar", t["ytutar"]), ("Vadesi gelen tutar", t["vadesi"]), ("Ödenen tutar", t["odenen"]), ("Eksik ödeme", t["eksik"]), ("Ödeme uyumu %", t["uyum"])):
        _satir(ws, [etiket, None, deger], para_s=(3,))
    ws.append([])
    ws.append([f"Durum tarihi: {bugun:%d.%m.%Y}"])
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, gecikme: int = 30, izleme_ay: int = 12) -> dict:
    krediler, uy = oku(yol)
    r = degerlendir(krediler, bugun, gecikme, izleme_ay)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, bugun)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Yeniden yapılandırılan kredilerin ödeme planı uyumunu ve izleme süresini izler.")
    p.add_argument("--krediler", type=Path, help="Müşteri No, Müşteri, Kredi No, Yapılandırma Tarihi, Yapılandırılan Tutar, Kalan Risk, Vadesi Gelen Tutar, Ödenen Tutar, Gecikme Gün, İzleme Süresi (ay), Yapılandırma Sayısı")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--gecikme", type=int, default=30, help="Gecikme bu günü aşarsa kredi 'Uyumsuz' sayılır (varsayılan 30)")
    p.add_argument("--izleme-ay", type=int, default=12, help="Satırda yazılmamışsa izleme süresi, ay (varsayılan 12)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "yapilandirma_takip_listesi.xlsx")
    a = p.parse_args(argv)
    ornek = a.krediler is None
    yol = ORNEK / "krediler.csv" if ornek else a.krediler
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    if bugun is None or a.gecikme < 0 or not 1 <= a.izleme_ay <= 120:
        print("[X] --tarih GG.AA.YYYY; --gecikme sıfır veya pozitif; --izleme-ay 1–120 arasında olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, bugun, a.gecikme, a.izleme_ay)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {len(r['krediler'])} yapılandırılmış kredi · kalan risk {t['risk']:,.2f} · " + ", ".join(f"{o['durum']} {o['kredi']}" for o in r["ozet"]) + (f" · ödeme uyumu %{t['uyum']}" if t["uyum"] is not None else ""))
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
