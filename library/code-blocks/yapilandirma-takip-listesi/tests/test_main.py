import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)


def k(no, gun=0, vadesi=100, odenen=None, ytarih=date(2026, 3, 1), izleme=None, sayi=1, risk=1000):
    return {"no": no, "ad": "Ad", "kredi": "K", "ytarih": ytarih, "ytutar": D(2000), "risk": D(str(risk)), "vadesi": D(str(vadesi)), "odenen": D(str(vadesi if odenen is None else odenen)), "gun": gun, "izleme": izleme, "sayi": sayi}


def d(krediler, **kw):
    r = main.degerlendir(krediler, BUGUN, **kw)
    return {x["no"]: x["durum"] for x in r["krediler"]}, sorted((u["onem"], u["tur"], u["kim"].split(" · ")[0]) for u in r["uyarilar"])


class DegerlendirTesti(unittest.TestCase):
    def test_durumlar(self):
        s, uy = d([k("0"), k("1", 1, odenen=90), k("30", 30, odenen=90), k("31", 31, odenen=90), k("eksik", 0, odenen="99.99"), k("fazla", 0, odenen=150), k("taksitsiz", 0, 0), k("doldu", ytarih=date(2025, 10, 12)),
                   k("yarın", ytarih=date(2025, 10, 13)), k("doldu_gecikmeli", 5, odenen=50, ytarih=date(2025, 1, 1)), k("iki", sayi=2)])
        self.assertEqual(s, {"0": "Uyumlu", "1": "Gecikmeli", "30": "Gecikmeli", "31": "Uyumsuz", "eksik": "Gecikmeli", "fazla": "Uyumlu", "taksitsiz": "Uyumlu", "doldu": "İzleme süresi doldu", "yarın": "Uyumlu",
                             "doldu_gecikmeli": "Gecikmeli", "iki": "Uyumlu"})
        self.assertEqual(uy, [("Bilgi", "İzleme süresi bitiyor", "yarın"), ("Bilgi", "İzleme süresi doldu", "doldu"), ("Orta", "Birden çok kez yapılandırılmış", "iki"), ("Orta", "Gecikme bilgisi tutarsız", "eksik"),
                              ("Orta", "Gecikmeli ödeme", "1"), ("Orta", "Gecikmeli ödeme", "30"), ("Orta", "Gecikmeli ödeme", "doldu_gecikmeli"), ("Orta", "Gecikmeli ödeme", "eksik"), ("Yüksek", "Ödeme planına uyumsuz", "31")])
        self.assertEqual((d([k("a", 16, odenen=1)], gecikme=15)[0]["a"], d([k("a", 15, odenen=1)], gecikme=15)[0]["a"], d([k("a", 1, odenen=1)], gecikme=0)[0]["a"]), ("Uyumsuz", "Gecikmeli", "Uyumsuz"))
        self.assertEqual([u[1] for u in d([k("a", ytarih=date(2026, 11, 1)), k("b", 3)])[1]], ["Gecikme bilgisi tutarsız", "Gecikmeli ödeme", "Yapılandırma tarihi ileride"])

    def test_hesaplar(self):
        l = [k("a", 45, 400, 250, risk=2300), k("b", 0, 0, 0, risk=700), k("c", 0, 300, 450, date(2026, 4, 12), 6, risk=1000), k("d", 0, 100, 100, date(2025, 8, 31), 6, risk=0), k("e", ytarih=date(2026, 4, 13), izleme=6, risk=0)]
        r = main.degerlendir(l, BUGUN, izleme_ay=24)
        x = {v["no"]: v for v in r["krediler"]}
        self.assertEqual([(x[n]["uyum"], x[n]["eksik"], x[n]["izleme_ay"], x[n]["izleme_bitis"], x[n]["izleme_kalan"], x[n]["durum"]) for n in "abcde"],
                         [(D("62.50"), D(150), 24, date(2028, 3, 1), 506, "Uyumsuz"), (None, D(0), 24, date(2028, 3, 1), 506, "Uyumlu"), (D("150.00"), D(0), 6, BUGUN, 0, "İzleme süresi doldu"),
                          (D("100.00"), D(0), 6, date(2026, 2, 28), -226, "İzleme süresi doldu"), (D("100.00"), D(0), 6, date(2026, 10, 13), 1, "Uyumlu")])
        self.assertEqual([v["no"] for v in r["krediler"]], ["a", "c", "d", "b", "e"])
        self.assertEqual(r["ozet"], [{"durum": "Uyumsuz", "kredi": 1, "risk": D(2300), "pay": D("57.50")}, {"durum": "Gecikmeli", "kredi": 0, "risk": D(0), "pay": D("0.00")}, {"durum": "İzleme süresi doldu", "kredi": 2, "risk": D(1000), "pay": D("25.00")},
                                     {"durum": "Uyumlu", "kredi": 2, "risk": D(700), "pay": D("17.50")}])
        self.assertEqual(r["toplam"], {"risk": D(4000), "ytutar": D(10000), "vadesi": D(900), "odenen": D(900), "eksik": D(150), "uyum": D("100.00")})
        t = main.degerlendir([k("a", 0, 0, 0, risk=0)], BUGUN)
        self.assertEqual((t["toplam"]["uyum"], t["ozet"][3]["pay"]), (None, None))
        with self.assertRaises(ValueError):
            main.degerlendir([], BUGUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "krediler.csv", cikti, BUGUN)
            self.assertEqual([(x["no"], x["durum"], x["uyum"], x["eksik"], x["izleme_bitis"], x["izleme_kalan"]) for x in r["krediler"]],
                             [("4002", "Uyumsuz", D("62.50"), D(150000), date(2027, 3, 5), 144), ("4004", "Gecikmeli", D("75.00"), D(20000), date(2027, 6, 15), 246), ("4006", "Gecikmeli", D("93.33"), D(10000), date(2026, 12, 12), 61),
                              ("4003", "İzleme süresi doldu", D("100.00"), D(0), date(2026, 9, 20), -22), ("4005", "Uyumlu", None, D(0), date(2027, 10, 1), 354), ("4001", "Uyumlu", D("100.00"), D(0), date(2027, 1, 10), 90)])
            self.assertEqual(r["ozet"], [{"durum": "Uyumsuz", "kredi": 1, "risk": D(2300000), "pay": D("38.59")}, {"durum": "Gecikmeli", "kredi": 2, "risk": D(860000), "pay": D("14.43")},
                                         {"durum": "İzleme süresi doldu", "kredi": 1, "risk": D(350000), "pay": D("5.87")}, {"durum": "Uyumlu", "kredi": 2, "risk": D(2450000), "pay": D("41.11")}])
            self.assertEqual(r["toplam"], {"risk": D(5960000), "ytutar": D(7000000), "vadesi": D(1430000), "odenen": D(1250000), "eksik": D(180000), "uyum": D("87.41")})
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Yüksek", "Ödeme planına uyumsuz", "4002 · Beta Yapı · Y-2"), ("Orta", "Birden çok kez yapılandırılmış", "4004 · Delta Tekstil · Y-4"), ("Orta", "Gecikme bilgisi tutarsız", "4006 · Zeta Mobilya · Y-6"),
                              ("Orta", "Gecikmeli ödeme", "4004 · Delta Tekstil · Y-4"), ("Orta", "Gecikmeli ödeme", "4006 · Zeta Mobilya · Y-6"), ("Bilgi", "İzleme süresi doldu", "4003 · Gama Gıda · Y-3")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Yapılandırmalar", "Özet", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "k.csv"
            y.write_text("Müşteri Numarası;Unvan;Hesap No;Yapılandırma;Yapılandırma Tutarı;Bakiye;Tahakkuk Eden;Tahsil Edilen;Gecikme;İzleme Ay;Kaçıncı Yapılandırma\n7;A;H1;01.03.2026;1.250.000,50;900.000;120.000;90.000;12;24;2\n"
                         "8;B;;01.03.2026;;;;;;;\n9;C;;x;;;;;;;\n10;C;;01.03.2026;-1;;;;;;\n11;C;;01.03.2026;;;;;1,5;;\n12;C;;01.03.2026;;;;;;0;\n;C;;01.03.2026;;;;;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual((l, [(u["tur"], u["kim"]) for u in uy]),
                             ([{"no": "7", "ad": "A", "kredi": "H1", "ytarih": date(2026, 3, 1), "ytutar": D("1250000.50"), "risk": D(900000), "vadesi": D(120000), "odenen": D(90000), "gun": 12, "izleme": 24, "sayi": 2},
                               {"no": "8", "ad": "B", "kredi": "", "ytarih": date(2026, 3, 1), "ytutar": D(0), "risk": D(0), "vadesi": D(0), "odenen": D(0), "gun": 0, "izleme": None, "sayi": 1}], [("Satır okunamadı", f"{n} · C") for n in (9, 10, 11, 12)]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--krediler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "12.10.2026", "--gecikme", "15", "--izleme-ay", "24"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--gecikme", "-1"], ["--izleme-ay", "0"])], [2] * 3)
            self.assertEqual(main.main(["--krediler", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Müşteri No;Kredi No\nA;B\n", "Müşteri No;Yapılandırma Tarihi\n;01.03.2026\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
