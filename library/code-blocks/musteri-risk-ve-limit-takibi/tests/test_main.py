import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)


def m(kod, limit=1000, teminat=0):
    return {"kod": kod, "ad": "Ad", "limit": None if limit is None else D(str(limit)), "teminat": D(teminat)}


def k(kod, tutar, vade=None):
    return {"kod": kod, "belge": "B", "vade": None if vade is None else BUGUN + timedelta(days=vade), "tutar": D(str(tutar))}


def s(kod, tutar, no="S"):
    return {"kod": kod, "no": no, "tutar": D(str(tutar))}


class KararTesti(unittest.TestCase):
    def test_sinirlar(self):
        mus = [m(x) for x in "ABCDEFGK"] + [m("H", None), m("I", None), m("L", 0)]
        kal = [k("A", 600, 5), k("B", 600, 5), k("C", 500), k("D", 500, 1), k("E", 100, -30), k("F", 100, -31), k("G", 100, 0), k("H", 10, 5), k("J", 5, 5), k("K", -200, -50), k("K", 100), k("L", 20, 5)]
        sip = [s("A", 400), s("B", "400.01"), s("C", 400), s("D", "399.9")]
        r = main.degerlendir(mus, kal, sip, BUGUN)
        self.assertEqual([(o["kod"], o["karar"]) for o in r["musteriler"]],
                         [("F", "Sevk etme"), ("B", "Sevk etme"), ("E", "Onay gerekli"), ("A", "Onay gerekli"), ("C", "Onay gerekli"), ("L", "Onay gerekli"), ("H", "Onay gerekli"), ("J", "Onay gerekli"),
                          ("D", "Uygun"), ("G", "Uygun"), ("I", "Uygun"), ("K", "Uygun")])
        o = {x["kod"]: x for x in r["musteriler"]}
        self.assertEqual([(o[x]["bakiye"], o[x]["siparis"], o[x]["sonrasi"], o[x]["kullanim"], o[x]["sonrasi_kullanim"], o[x]["kullanilabilir"]) for x in "ABDK"],
                         [(D(600), D(400), D(1000), D("60.00"), D("100.00"), D(0)), (D(600), D("400.01"), D("1000.01"), D("60.00"), D("100.00"), D("-0.01")), (D(500), D("399.9"), D("899.9"), D("50.00"), D("89.99"), D("100.1")),
                          (D(-100), D(0), D(-100), D("-10.00"), D("-10.00"), D(1100))])
        self.assertEqual([(o[x]["vadesi_gecmis"], o[x]["en_eski"]) for x in "EFGK"], [(D(100), 30), (D(100), 31), (D(0), 0), (D(0), 0)])
        self.assertEqual([(o[x]["limit"], o[x]["kullanim"], o[x]["kullanilabilir"]) for x in "HIL"], [(None, None, None)] * 3)
        self.assertEqual((o["B"]["gerekce"], o["F"]["gerekce"], o["E"]["gerekce"], o["A"]["gerekce"], o["H"]["gerekce"], o["J"]["gerekce"], o["D"]["gerekce"]),
                         ("Sipariş sonrası bakiye limiti 0.01 aşıyor", "31 gündür vadesi geçmiş alacak var", "Vadesi geçmiş 100.00 alacak var (30 gün)", "Sipariş sonrası limit kullanımı %100.00", "Kredi limiti tanımlı değil",
                          "Müşteri dosyasında tanımlı değil", ""))
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         sorted([("Yüksek", "Limit aşımı", "B · Ad"), ("Yüksek", "Uzun süredir vadesi geçmiş alacak", "F · Ad"), ("Orta", "Vadesi geçmiş alacak", "E · Ad"), ("Orta", "Limite yaklaştı", "A · Ad"), ("Orta", "Limite yaklaştı", "C · Ad"),
                                 ("Orta", "Limit tanımsız", "H · Ad"), ("Orta", "Limit tanımsız", "L · Ad"), ("Orta", "Tanımsız müşteri kodu", "J"), ("Bilgi", "Alacak bakiyesi", "K · Ad")]))
        self.assertEqual((r["say"], r["bakiye"], r["vadesi_gecmis"], r["siparis"], r["siparis_karar"]),
                         ({"Sevk etme": 2, "Onay gerekli": 6, "Uygun": 4}, D(2435), D(200), D("1599.91"), {"Sevk etme": D("400.01"), "Onay gerekli": D(800), "Uygun": D("399.9")}))
        self.assertEqual([(x["kod"], x["karar"]) for x in r["siparisler"]], [("B", "Sevk etme"), ("A", "Onay gerekli"), ("C", "Onay gerekli"), ("D", "Uygun")])

    def test_esikler_ve_ozel_durumlar(self):
        kr = lambda sip, vade, **kw: main.degerlendir([m("A")], [k("A", 500, vade)], [s("A", sip)], BUGUN, **kw)["musteriler"][0]["karar"]   # noqa: E731
        self.assertEqual([kr(300, 5, esik=D(80)), kr("299.9", 5, esik=D(80)), kr(0, -15, gecikme=15), kr(0, -16, gecikme=15), kr(0, -1, gecikme=0), kr(600, -40)],
                         ["Onay gerekli", "Uygun", "Onay gerekli", "Sevk etme", "Sevk etme", "Sevk etme"])
        r = main.degerlendir([m("A", 1000, 300), m("a", 5), m("B", 1000, 5000)], [k("A", 400, -40), k("A", 800, 3), k("B", 700, 3)], [s("A", 100, "S2"), s("a", 50, "S1")], BUGUN)
        o = r["musteriler"][0]
        self.assertEqual((o["kod"], o["bakiye"], o["siparis"], o["siparis_adet"], o["sonrasi"], o["sonrasi_kullanim"], o["teminatsiz"], o["karar"], o["gerekce"], r["musteriler"][1]["teminatsiz"]),
                         ("A", D(1200), D(150), 2, D(1350), D("135.00"), D(900), "Sevk etme", "Sipariş sonrası bakiye limiti 350.00 aşıyor; 40 gündür vadesi geçmiş alacak var", D(0)))
        self.assertEqual(([x["no"] for x in r["siparisler"]], sorted(u["tur"] for u in r["uyarilar"])), (["S1", "S2"], ["Limit aşımı", "Mükerrer müşteri kodu", "Uzun süredir vadesi geçmiş alacak"]))
        r = main.degerlendir([m("A")], [k("A", 100, 5)], [], BUGUN)
        self.assertEqual((r["musteriler"][0]["karar"], r["siparisler"], r["siparis"]), ("Uygun", [], D(0)))
        with self.assertRaises(ValueError):
            main.degerlendir([], [], [], BUGUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "musteriler.csv", ORNEK / "acik_kalemler.csv", ORNEK / "siparisler.csv", cikti, BUGUN)
            self.assertEqual([(o["kod"], o["karar"], o["bakiye"], o["kullanim"], o["vadesi_gecmis"], o["en_eski"], o["siparis"], o["sonrasi"], o["sonrasi_kullanim"], o["kullanilabilir"], o["teminatsiz"]) for o in r["musteriler"]],
                             [("M002", "Sevk etme", D(150000), D("50.00"), D(90000), 53, D(50000), D(200000), D("66.67"), D(100000), D(150000)),
                              ("M004", "Sevk etme", D(340000), D("85.00"), D(0), 0, D(80000), D(420000), D("105.00"), D(-20000), D(340000)),
                              ("M003", "Onay gerekli", D(180000), D("72.00"), D(70000), 7, D(40000), D(220000), D("88.00"), D(30000), D(80000)),
                              ("M001", "Onay gerekli", D(300000), D("60.00"), D(0), 0, D(150000), D(450000), D("90.00"), D(50000), D(100000)),
                              ("M005", "Onay gerekli", D(45000), None, D(0), 0, D(30000), D(75000), None, None, D(45000)), ("M009", "Onay gerekli", D(0), None, D(0), 0, D(20000), D(20000), None, None, D(0)),
                              ("M006", "Uygun", D(25000), D("16.67"), D(0), 0, D(60000), D(85000), D("56.67"), D(65000), D(25000))])
            self.assertEqual((r["say"], r["bakiye"], r["vadesi_gecmis"], r["siparis"], r["siparis_karar"]),
                             ({"Sevk etme": 2, "Onay gerekli": 4, "Uygun": 1}, D(1040000), D(160000), D(430000), {"Sevk etme": D(130000), "Onay gerekli": D(240000), "Uygun": D(60000)}))
            self.assertEqual([(x["no"], x["karar"]) for x in r["siparisler"]],
                             [("S-502", "Sevk etme"), ("S-504", "Sevk etme"), ("S-501", "Onay gerekli"), ("S-503", "Onay gerekli"), ("S-505", "Onay gerekli"), ("S-508", "Onay gerekli"), ("S-506", "Uygun")])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Yüksek", "Limit aşımı", "M004 · Delta Tekstil"), ("Yüksek", "Uzun süredir vadesi geçmiş alacak", "M002 · Beta Yapı"), ("Orta", "Limit tanımsız", "M005 · Epsilon Otomotiv"),
                              ("Orta", "Limite yaklaştı", "M001 · Alfa Market"), ("Orta", "Tanımsız müşteri kodu", "M009"), ("Orta", "Vadesi geçmiş alacak", "M003 · Gama Gıda")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Müşteriler", "Siparişler", "Uyarılar"])
            r = main.calistir(ORNEK / "musteriler.csv", ORNEK / "acik_kalemler.csv", None, cikti, BUGUN)
            self.assertEqual((r["say"], load_workbook(cikti).sheetnames), ({"Sevk etme": 1, "Onay gerekli": 2, "Uygun": 3}, ["Müşteriler", "Uyarılar"]))

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            my, ky, sy = Path(tmp) / "m.csv", Path(tmp) / "k.csv", Path(tmp) / "s.csv"
            my.write_text("Cari Kodu;Unvan;Limit;Teminat Tutarı\nC1;A Ltd.;1.250.000;250.000,50\nC2;B;;\nC3;C;x;\n;D;5;\n", encoding="utf-8")
            ky.write_text("Cari Kodu;Fatura No;Vade Tarihi;Kalan Tutar\nC1;F1;15.10.2026;1.000,25\nC1;F2;;-50\nC2;F3;x;x\n", encoding="utf-8")
            sy.write_text("Kod;No;Toplam\nC1;S1;500\nC1;S2;-1\n", encoding="utf-8")
            l, uy = main.oku_musteriler(my)
            self.assertEqual((l, [(u["tur"], u["kim"]) for u in uy]),
                             ([{"kod": "C1", "ad": "A Ltd.", "limit": D(1250000), "teminat": D("250000.50")}, {"kod": "C2", "ad": "B", "limit": None, "teminat": D(0)}], [("Müşteri okunamadı", "C3")]))
            l, uy = main.oku_kalemler(ky)
            self.assertEqual(([(x["belge"], x["vade"], x["tutar"]) for x in l], [u["tur"] for u in uy]), ([("F1", date(2026, 10, 15), D("1000.25")), ("F2", None, D(-50))], ["Açık kalem okunamadı"]))
            l, uy = main.oku_siparisler(sy)
            self.assertEqual((l, [u["tur"] for u in uy]), ([{"kod": "C1", "no": "S1", "tutar": D(500)}], ["Sipariş okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--musteriler", str(my), "--kalemler", str(ky), "--cikti", c]
            self.assertEqual(main.main(arg + ["--siparisler", str(sy), "--tarih", "12.10.2026", "--esik", "80", "--gecikme", "15"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + x) for x in (["--tarih", "x"], ["--esik", "0"], ["--esik", "101"], ["--gecikme", "-1"])] + [main.main(["--musteriler", str(my)])], [2] * 5)
            self.assertEqual(main.main(arg + ["--siparisler", str(Path(tmp) / "yok.csv")]), 1)
            ky.write_text("Müşteri Kodu;Vade\nA;1\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            my.write_text("Müşteri Kodu;Müşteri\n;\n", encoding="utf-8")
            ky.write_text("Müşteri Kodu;Tutar\n;\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
