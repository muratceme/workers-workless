"""
Müşteri Risk ve Limit Takibi — Gerisi Bende kod bloğu
Satış › Satış Destek Uzmanı

Açık bakiye, vadesi geçmiş alacak ve kredi limitine bakarak sevkiyat öncesi riskli müşterileri işaretler.
  Açık bakiye            = müşterinin açık kalemlerinin toplamı   (eksi kalemler: avans, iade, alacak dekontu)
  Vadesi geçmiş alacak   = vadesi durum tarihinden önce olan artı kalemlerin toplamı
  En eski gecikme (gün)  = durum tarihi − en eski vadesi geçmiş kalemin vadesi
  Sipariş sonrası bakiye = açık bakiye + sevk bekleyen sipariş tutarı
  Limit kullanımı %      = bakiye ÷ kredi limiti × 100   (sipariş sonrası için de ayrıca hesaplanır)
  Karar                  = Sevk etme     : sipariş sonrası bakiye limiti aşıyor ya da gecikme 30 günden uzun
                           Onay gerekli  : vadesi geçmiş alacak var, limit tanımsız ya da sipariş sonrası kullanım %90 ve üstünde
                           Uygun         : hiçbiri yok
Karar kuralı bir çalışma varsayımıdır; eşikler komut satırından değişir. Son kararı kredi politikanız ve yetkiliniz verir.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
M_SUTUN = {"kod": ("musteri kodu", "cari kodu", "kod"), "ad": ("musteri", "musteri adi", "cari adi", "unvan", "firma"), "limit": ("kredi limiti", "limit", "risk limiti"), "teminat": ("teminat", "teminat tutari")}
K_SUTUN = {"kod": ("musteri kodu", "cari kodu", "kod"), "belge": ("belge no", "fatura no", "evrak no"), "vade": ("vade", "vade tarihi"), "tutar": ("tutar", "kalan tutar", "bakiye", "acik tutar")}
S_SUTUN = {"kod": ("musteri kodu", "cari kodu", "kod"), "no": ("siparis no", "siparis numarasi", "no"), "tutar": ("tutar", "siparis tutari", "toplam")}
KARAR_SIRA = {"Sevk etme": 0, "Onay gerekli": 1, "Uygun": 2}
KARAR_RENK = {"Sevk etme": KIRMIZI, "Onay gerekli": SARI, "Uygun": YESIL}
Z, YUZ = Decimal(0), Decimal(100)


def _oku(yol: Path, sozluk: dict, zorunlu: tuple, ne: str, kur) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, sozluk, zorunlu):
        kod = metin(r["kod"])
        if not kod:
            continue
        kayit = kur(r, kod)
        if kayit is None:
            uy.append(uyari("Yüksek", f"{ne} okunamadı", kod, f"{yol.name} satır {r['_satir']}: tutar okunamadı ya da geçersiz. Satır alınmadı."))
        else:
            l.append(kayit)
    return l, uy


def oku_musteriler(yol: Path) -> tuple[list[dict], list[dict]]:
    def kur(r, kod):
        limit, teminat = para(r["limit"]), para(r["teminat"])
        if (metin(r["limit"]) and (limit is None or limit < 0)) or (metin(r["teminat"]) and (teminat is None or teminat < 0)):
            return None
        return {"kod": kod, "ad": metin(r["ad"]), "limit": limit, "teminat": teminat or Z}
    return _oku(yol, M_SUTUN, ("kod",), "Müşteri", kur)


def oku_kalemler(yol: Path) -> tuple[list[dict], list[dict]]:
    def kur(r, kod):
        tutar = para(r["tutar"])
        return None if tutar is None else {"kod": kod, "belge": metin(r["belge"]), "vade": tarih(r["vade"]), "tutar": tutar}
    return _oku(yol, K_SUTUN, ("kod", "tutar"), "Açık kalem", kur)


def oku_siparisler(yol: Path) -> tuple[list[dict], list[dict]]:
    def kur(r, kod):
        tutar = para(r["tutar"])
        return None if tutar is None or tutar < 0 else {"kod": kod, "no": metin(r["no"]), "tutar": tutar}
    return _oku(yol, S_SUTUN, ("kod", "tutar"), "Sipariş", kur)


def degerlendir(musteriler: list[dict], kalemler: list[dict], siparisler: list[dict], bugun: date, esik: Decimal = Decimal(90), gecikme: int = 30) -> dict:
    if not musteriler and not kalemler and not siparisler:
        raise ValueError("Müşteri, açık kalem ve sipariş listesi boş.")
    uy, m = [], {}
    for x in musteriler:
        if katla(x["kod"]) in m:
            uy.append(uyari("Orta", "Mükerrer müşteri kodu", x["kod"], "Müşteri dosyasında birden çok kez geçiyor; ilk satır kullanıldı."))
            continue
        m[katla(x["kod"])] = {**x, "tanimli": True, "kalemler": [], "siparisler": []}
    for liste, alan in ((kalemler, "kalemler"), (siparisler, "siparisler")):
        for x in liste:
            m.setdefault(katla(x["kod"]), {"kod": x["kod"], "ad": "", "limit": None, "teminat": Z, "tanimli": False, "kalemler": [], "siparisler": []})[alan].append(x)
    sonuc = []
    for c in m.values():
        kim = f"{c['kod']} · {c['ad']}".strip(" ·")
        gecmis = [k for k in c["kalemler"] if k["tutar"] > 0 and k["vade"] and k["vade"] < bugun]
        o = {"kod": c["kod"], "ad": c["ad"], "limit": c["limit"] or None, "teminat": c["teminat"], "bakiye": sum((k["tutar"] for k in c["kalemler"]), Z), "vadesi_gecmis": sum((k["tutar"] for k in gecmis), Z),
             "en_eski": max(((bugun - k["vade"]).days for k in gecmis), default=0), "siparis": sum((s["tutar"] for s in c["siparisler"]), Z), "siparis_adet": len(c["siparisler"])}
        o["sonrasi"] = o["bakiye"] + o["siparis"]
        o["kullanim"] = yuvarla(o["bakiye"] * YUZ / o["limit"]) if o["limit"] else None
        o["sonrasi_kullanim"] = yuvarla(o["sonrasi"] * YUZ / o["limit"]) if o["limit"] else None
        o["kullanilabilir"] = o["limit"] - o["sonrasi"] if o["limit"] else None
        o["teminatsiz"] = max(o["bakiye"] - o["teminat"], Z)
        engel, onay = [], []
        if o["limit"] and o["sonrasi"] > o["limit"]:
            engel.append(f"sipariş sonrası bakiye limiti {o['sonrasi'] - o['limit']:,.2f} aşıyor")
            uy.append(uyari("Yüksek", "Limit aşımı", kim, f"Limit {o['limit']:,.2f}; açık bakiye {o['bakiye']:,.2f}, bekleyen sipariş {o['siparis']:,.2f}; sipariş sonrası {o['sonrasi']:,.2f} (%{o['sonrasi_kullanim']})."))
        if o["en_eski"] > gecikme:
            engel.append(f"{o['en_eski']} gündür vadesi geçmiş alacak var")
            uy.append(uyari("Yüksek", "Uzun süredir vadesi geçmiş alacak", kim, f"Vadesi geçmiş {o['vadesi_gecmis']:,.2f}; en eski kalem {o['en_eski']} gün gecikmiş (eşik {gecikme} gün)."))
        elif o["vadesi_gecmis"] > 0:
            onay.append(f"vadesi geçmiş {o['vadesi_gecmis']:,.2f} alacak var ({o['en_eski']} gün)")
            uy.append(uyari("Orta", "Vadesi geçmiş alacak", kim, f"Vadesi geçmiş {o['vadesi_gecmis']:,.2f}; en eski kalem {o['en_eski']} gün gecikmiş."))
        if not c["tanimli"]:
            onay.append("müşteri dosyasında tanımlı değil")
            uy.append(uyari("Orta", "Tanımsız müşteri kodu", kim, "Açık kalem ya da siparişte geçiyor ama müşteri dosyasında yok; limit bilinmiyor."))
        elif not o["limit"] and o["sonrasi"] > 0:
            onay.append("kredi limiti tanımlı değil")
            uy.append(uyari("Orta", "Limit tanımsız", kim, f"Kredi limiti yazılmamış; açık bakiye {o['bakiye']:,.2f}, bekleyen sipariş {o['siparis']:,.2f}."))
        if o["limit"] and o["sonrasi"] <= o["limit"] and o["sonrasi_kullanim"] >= esik:
            onay.append(f"sipariş sonrası limit kullanımı %{o['sonrasi_kullanim']}")
            uy.append(uyari("Orta", "Limite yaklaştı", kim, f"Sipariş sonrası bakiye {o['sonrasi']:,.2f}, limit {o['limit']:,.2f} (%{o['sonrasi_kullanim']}); kalan {o['kullanilabilir']:,.2f}."))
        if o["bakiye"] < 0:
            uy.append(uyari("Bilgi", "Alacak bakiyesi", kim, f"Müşterinin bakiyesi {o['bakiye']:,.2f}; avans ya da iade kaydı fazla olabilir."))
        o["karar"] = "Sevk etme" if engel else "Onay gerekli" if onay else "Uygun"
        o["gerekce"] = "; ".join(engel + onay).capitalize() if engel or onay else ""
        sonuc.append(o)
    sonuc.sort(key=lambda o: (KARAR_SIRA[o["karar"]], -o["vadesi_gecmis"], -o["bakiye"], katla(o["kod"])))
    karar = {katla(o["kod"]): o for o in sonuc}
    sip = [{**s, "ad": karar[katla(s["kod"])]["ad"], "karar": karar[katla(s["kod"])]["karar"], "gerekce": karar[katla(s["kod"])]["gerekce"]} for s in siparisler]
    sip.sort(key=lambda s: (KARAR_SIRA[s["karar"]], katla(s["kod"]), katla(s["no"])))
    return {"musteriler": sonuc, "siparisler": sip, "say": {k: sum(1 for o in sonuc if o["karar"] == k) for k in KARAR_SIRA}, "bakiye": sum((o["bakiye"] for o in sonuc), Z), "vadesi_gecmis": sum((o["vadesi_gecmis"] for o in sonuc), Z),
            "siparis": sum((o["siparis"] for o in sonuc), Z), "siparis_karar": {k: sum((s["tutar"] for s in sip if s["karar"] == k), Z) for k in KARAR_SIRA}, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Müşteriler"
    _baslik(ws, ["Karar", "Müşteri Kodu", "Müşteri", "Kredi Limiti", "Açık Bakiye", "Kullanım %", "Vadesi Geçmiş", "En Eski Gecikme (gün)", "Bekleyen Sipariş", "Sipariş Sonrası Bakiye", "Sipariş Sonrası Kullanım %", "Kalan Limit",
                 "Teminat", "Teminatsız Risk", "Gerekçe"], (14, 13, 28, 15, 15, 10, 15, 10, 15, 15, 11, 15, 14, 15, 70))
    for o in r["musteriler"]:
        _satir(ws, [o["karar"], o["kod"], o["ad"], o["limit"], o["bakiye"], o["kullanim"], o["vadesi_gecmis"] or None, o["en_eski"] or None, o["siparis"] or None, o["sonrasi"], o["sonrasi_kullanim"], o["kullanilabilir"],
                    o["teminat"] or None, o["teminatsiz"], o["gerekce"]], para_s=(4, 5, 6, 7, 9, 10, 11, 12, 13, 14), dolgu={1: KARAR_RENK[o["karar"]], 12: KIRMIZI if o["kullanilabilir"] is not None and o["kullanilabilir"] < 0 else None}, sar=(15,))
    if r["siparisler"]:
        ws = wb.create_sheet("Siparişler")
        _baslik(ws, ["Karar", "Sipariş No", "Müşteri Kodu", "Müşteri", "Sipariş Tutarı", "Gerekçe"], (14, 14, 13, 28, 15, 70))
        for s in r["siparisler"]:
            _satir(ws, [s["karar"], s["no"], s["kod"], s["ad"], s["tutar"], s["gerekce"]], para_s=(5,), dolgu={1: KARAR_RENK[s["karar"]]}, sar=(6,))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(m_yol: Path, k_yol: Path, s_yol: Path | None, cikti: Path, bugun: date, esik: Decimal = Decimal(90), gecikme: int = 30) -> dict:
    m, uy1 = oku_musteriler(m_yol)
    k, uy2 = oku_kalemler(k_yol)
    s, uy3 = oku_siparisler(s_yol) if s_yol else ([], [])
    r = degerlendir(m, k, s, bugun, esik, gecikme)
    r["uyarilar"] = sirala(uy1 + uy2 + uy3 + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Açık bakiye, vadesi geçmiş alacak ve kredi limitine göre sevkiyat öncesi riskli müşterileri işaretler.")
    p.add_argument("--musteriler", type=Path, help="Müşteri Kodu, Müşteri, Kredi Limiti, Teminat")
    p.add_argument("--kalemler", type=Path, help="Açık kalemler: Müşteri Kodu, Belge No, Vade, Tutar")
    p.add_argument("--siparisler", type=Path, help="Sevk bekleyen siparişler: Müşteri Kodu, Sipariş No, Tutar (isteğe bağlı)")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--esik", default="90", help="Sipariş sonrası limit kullanımı bu yüzde ve üstündeyse onay istenir (varsayılan 90)")
    p.add_argument("--gecikme", type=int, default=30, help="En eski gecikme bu günden uzunsa sevk durdurulur (varsayılan 30)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "musteri_risk_limit_takibi.xlsx")
    a = p.parse_args(argv)
    if (a.musteriler is None) != (a.kalemler is None):
        print("[X] --musteriler ve --kalemler birlikte verilmelidir")
        return 2
    ornek = a.musteriler is None
    m_yol, k_yol = (ORNEK / "musteriler.csv", ORNEK / "acik_kalemler.csv") if ornek else (a.musteriler, a.kalemler)
    s_yol = ORNEK / "siparisler.csv" if ornek and a.siparisler is None else a.siparisler
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    esik = para(a.esik)
    if bugun is None or esik is None or not 0 < esik <= 100 or a.gecikme < 0:
        print("[X] --tarih GG.AA.YYYY; --esik 0–100 arasında; --gecikme sıfır veya pozitif olmalı")
        return 2
    for y in (m_yol, k_yol, s_yol):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(m_yol, k_yol, s_yol, a.cikti, bugun, esik, a.gecikme)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    s = r["say"]
    print(f"[OK] {len(r['musteriler'])} müşteri · açık bakiye {r['bakiye']:,.2f} · vadesi geçmiş {r['vadesi_gecmis']:,.2f} · Sevk etme {s['Sevk etme']}, Onay gerekli {s['Onay gerekli']}, Uygun {s['Uygun']}"
          + (f" · durdurulan sipariş {r['siparis_karar']['Sevk etme']:,.2f}" if r["siparisler"] else ""))
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
