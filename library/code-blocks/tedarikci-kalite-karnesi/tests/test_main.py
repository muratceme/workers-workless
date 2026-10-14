import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def p(ted, gelen, ret=0, karar=None, malzeme="M", t=date(2026, 7, 1)):
    return {"tarih": t, "tedarikci": ted, "malzeme": malzeme, "parti": "", "gelen": D(gelen), "ret": D(ret), "karar": karar or ("Ret" if ret == gelen else "Kabul")}


def veri():
    return [p("T1", 1000), p("t1", 1000, 2), p("T1", 2000, 2, malzeme="N", t=date(2026, 8, 1)), p("T2", 500, 10), p("T2", 500), p("T2", 1000, t=date(2026, 8, 1)), p("T3", 100, 1), p("T3", 100, 1), p("T3", 100, 1, "Şartlı kabul"),
            p("T4", 100, 100, t=None)]


def k(g):
    return (g["parti"], g["gelen"], g["ret"], g["ppm"], g["ret_oran"], g["hatali"], g["hatali_oran"], g["reddedilen"], g["sartli"])


class KarneTesti(unittest.TestCase):
    def test_notlar(self):
        r = main.karne(veri())
        self.assertEqual([(t["tedarikci"], t["karne"]) + k(t) for t in r["tedarikciler"]],
                         [("T3", "C", 3, D(300), D(3), D(10000), D("1.000"), 3, D("100.00"), 0, 1), ("T2", "B", 3, D(2000), D(10), D(5000), D("0.500"), 1, D("33.33"), 0, 0),
                          ("T1", "A", 3, D(4000), D(4), D(1000), D("0.100"), 2, D("66.67"), 0, 0), ("T4", "Az veri", 1, D(100), D(100), D(1000000), D("100.000"), 1, D("100.00"), 1, 0)])
        self.assertEqual(k(r["toplam"]), (10, D(6400), D(117), D(18281), D("1.828"), 7, D("70.00"), 1, 1))
        self.assertEqual([(d["tedarikci"], d["malzeme"], d["parti"], d["ppm"]) for d in r["detay"]], [("T1", "M", 2, D(1000)), ("T1", "N", 1, D(1000)), ("T2", "M", 3, D(5000)), ("T3", "M", 3, D(10000)), ("T4", "M", 1, D(1000000))])
        self.assertEqual([(a["ay"], a["parti"], a["gelen"], a["ret"], a["ppm"]) for a in r["aylar"]], [("2026-07", 7, D(3300), D(15), D(4545)), ("2026-08", 2, D(3000), D(2), D(667))])
        self.assertEqual((r["notlar"], sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"])),
                         ({"A": 1, "B": 1, "C": 1, "Az veri": 1}, [("Bilgi", "Az veride ret var", "T4"), ("Orta", "Karne notu B", "T2"), ("Yüksek", "Karne notu C", "T3")]))

    def test_esikler(self):
        n = lambda **kw: {t["tedarikci"]: t["karne"] for t in main.karne(veri(), **kw)["tedarikciler"]}   # noqa: E731
        self.assertEqual((n(a=D(999), b=D(4999)), n(asgari=1), n(asgari=4), n(a=D(0), b=D(10000))),
                         ({"T1": "B", "T2": "C", "T3": "C", "T4": "Az veri"}, {"T1": "A", "T2": "B", "T3": "C", "T4": "C"}, dict.fromkeys(("T1", "T2", "T3", "T4"), "Az veri"), {"T1": "B", "T2": "B", "T3": "B", "T4": "Az veri"}))
        with self.assertRaises(ValueError):
            main.karne([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "giris_kontrol.csv", cikti)
            self.assertEqual([(t["tedarikci"], t["karne"], t["parti"], t["gelen"], t["ret"], t["ppm"], t["hatali"], t["reddedilen"], t["sartli"]) for t in r["tedarikciler"]],
                             [("Numune Döküm San.", "C", 3, D(3000), D(1220), D(406667), 2, 1, 0), ("Deneme Plastik Ltd. Şti.", "B", 4, D(20000), D(50), D(2500), 3, 0, 1), ("Örnek Metal A.Ş.", "A", 5, D(120000), D(40), D(333), 3, 0, 0),
                              ("Kurgu Kimya Ltd. Şti.", "Az veri", 2, D(500), D(0), D(0), 0, 0, 0)])
            t = r["toplam"]
            self.assertEqual((t["parti"], t["gelen"], t["ret"], t["ppm"], t["hatali"]), (14, D(143500), D(1310), D(9129), 8))
            self.assertEqual([(d["tedarikci"], d["malzeme"], d["ppm"]) for d in r["detay"] if d["tedarikci"].startswith("Örnek")], [("Örnek Metal A.Ş.", "Cıvata M8", D(378)), ("Örnek Metal A.Ş.", "Sac 2 mm", D(200))])
            self.assertEqual([(a["ay"], a["parti"], a["ret"]) for a in r["aylar"]], [("2026-07", 4, D(35)), ("2026-08", 4, D(1206)), ("2026-09", 6, D(69))])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]], [("Yüksek", "Karne notu C", "Numune Döküm San."), ("Orta", "Karne notu B", "Deneme Plastik Ltd. Şti.")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Karne", "Tedarikçi × Malzeme", "Aylar", "Partiler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "g.csv"
            y.write_text("Giriş Tarihi;Firma;Stok Kodu;Lot No;Miktar;Red Miktarı;Sonuç\n01.07.2026;A;S1;L1;1.000;5;ŞARTLI\n;A;;;500;500;\n;A;;L3;500;;Bekliyor\n;A;;L4;10;11;\n;A;;L5;0;;\n;A;;L6;x;;\n;;;;10;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["tarih"], x["tedarikci"], x["malzeme"], x["parti"], x["gelen"], x["ret"], x["karar"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([(date(2026, 7, 1), "A", "S1", "L1", D(1000), D(5), "Şartlı kabul"), (None, "A", "(malzeme yok)", "", D(500), D(500), "Ret"), (None, "A", "(malzeme yok)", "L3", D(500), D(0), "Kabul")],
                              [("Karar tanınmadı", "A · L3"), ("Parti okunamadı", "A · L4"), ("Parti okunamadı", "A · L5"), ("Parti okunamadı", "A · L6")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--kayitlar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--a", "500", "--b", "2.500", "--min", "1"]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--a", "5000"], ["--a", "x"], ["--b", "-1"], ["--min", "0"])], [2] * 4)
            self.assertEqual(main.main(["--kayitlar", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Tedarikçi;Malzeme\nA;B\n", "Tedarikçi;Gelen Miktar\n;5\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
