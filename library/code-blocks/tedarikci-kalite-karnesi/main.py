"""
Tedarikçi Kalite Karnesi — Gerisi Bende kod bloğu
Kalite Güvence ve Kalite Kontrol › Kalite Müdürü

Giriş kalite kontrol kayıtlarından tedarikçi bazında ret oranını (ppm) hesaplar ve karne notu verir.
  Ret oranı (ppm)        = Σ ret miktarı ÷ Σ gelen miktar × 1.000.000
  Ret oranı %            = Σ ret miktarı ÷ Σ gelen miktar × 100
  Hatalı parti oranı %   = ret miktarı sıfırdan büyük ya da kararı Ret / Şartlı kabul olan parti ÷ toplam parti × 100
  Reddedilen parti       = kararı "Ret" olan (karar yazılmamışsa tamamı reddedilen) parti
  Karne notu             = A (ppm ≤ --a, 1.000) · B (ppm ≤ --b, 5.000) · C (üstü);   parti sayısı --min (3) altındaysa "Az veri"
Not eşikleri sizin politikanızdır; ürün ve sektöre göre değişir. Teslimat ve fiyat performansı bu karnede yoktur.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"tarih": ("tarih", "giris tarihi", "kontrol tarihi"), "tedarikci": ("tedarikci", "tedarikci adi", "firma"), "malzeme": ("malzeme", "malzeme adi", "stok kodu", "urun"), "parti": ("parti no", "parti", "irsaliye no", "lot no"),
         "gelen": ("gelen miktar", "giris miktari", "teslim alinan miktar", "miktar"), "ret": ("ret miktar", "ret miktari", "red miktari", "uygun olmayan miktar"), "karar": ("karar", "sonuc", "kabul durumu")}
KARAR = {"kabul": "Kabul", "ret": "Ret", "red": "Ret", "sartli kabul": "Şartlı kabul", "sartli": "Şartlı kabul"}
NOT_RENK = {"A": YESIL, "B": SARI, "C": KIRMIZI, "Az veri": None}
MILYON = Decimal(1000000)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("tedarikci", "gelen")):
        ted = metin(r["tedarikci"])
        if not ted:
            continue
        kim = f"{ted} · {metin(r['parti']) or 'satır ' + str(r['_satir'])}"
        gelen, ret = para(r["gelen"]), para(r["ret"]) if metin(r["ret"]) else Decimal(0)
        if gelen is None or gelen <= 0 or ret is None or not 0 <= ret <= gelen:
            uy.append(uyari("Yüksek", "Parti okunamadı", kim, f"Satır {r['_satir']}: gelen miktar '{metin(r['gelen'])}' ya da ret miktarı '{metin(r['ret'])}' okunamadı veya ret gelenden büyük. Satır alınmadı."))
            continue
        karar = KARAR.get(katla(r["karar"]))
        if karar is None and metin(r["karar"]):
            uy.append(uyari("Bilgi", "Karar tanınmadı", kim, f"'{metin(r['karar'])}' Kabul, Ret ya da Şartlı kabul olarak okunamadı; ret miktarına göre değerlendirildi."))
        l.append({"tarih": tarih(r["tarih"]), "tedarikci": ted, "malzeme": metin(r["malzeme"]) or "(malzeme yok)", "parti": metin(r["parti"]), "gelen": gelen, "ret": ret,
                  "karar": karar or ("Ret" if ret == gelen else "Kabul")})
    return l, uy


def gosterge(l: list[dict]) -> dict:
    gelen, ret = sum((p["gelen"] for p in l), Decimal(0)), sum((p["ret"] for p in l), Decimal(0))
    hatali = sum(1 for p in l if p["ret"] > 0 or p["karar"] != "Kabul")
    return {"parti": len(l), "gelen": gelen, "ret": ret, "ppm": yuvarla(ret * MILYON / gelen, "1"), "ret_oran": yuvarla(ret * 100 / gelen, "0.001"), "hatali": hatali, "hatali_oran": yuvarla(Decimal(hatali * 100) / len(l)),
            "reddedilen": sum(1 for p in l if p["karar"] == "Ret"), "sartli": sum(1 for p in l if p["karar"] == "Şartlı kabul")}


def karne(partiler: list[dict], a: Decimal = Decimal(1000), b: Decimal = Decimal(5000), asgari: int = 3) -> dict:
    if not partiler:
        raise ValueError("Değerlendirilecek giriş kontrol kaydı yok.")

    def kirilim(*alanlar: str) -> list[dict]:
        grup = {}
        for p in partiler:
            grup.setdefault(tuple(katla(p[x]) for x in alanlar), []).append(p)
        return [{**{x: g[0][x] for x in alanlar}, **gosterge(g)} for g in grup.values()]

    tedarikciler, uy = kirilim("tedarikci"), []
    for t in tedarikciler:
        t["karne"] = "Az veri" if t["parti"] < asgari else "A" if t["ppm"] <= a else "B" if t["ppm"] <= b else "C"
        bilgi = f"{t['parti']} partide gelen {t['gelen']:,.0f}, ret {t['ret']:,.0f} ({t['ppm']:,.0f} ppm); hatalı parti {t['hatali']}, reddedilen parti {t['reddedilen']}."
        if t["karne"] == "C":
            uy.append(uyari("Yüksek", "Karne notu C", t["tedarikci"], f"{bilgi} Eşik {b:,.0f} ppm."))
        elif t["karne"] == "B":
            uy.append(uyari("Orta", "Karne notu B", t["tedarikci"], f"{bilgi} A notu için eşik {a:,.0f} ppm."))
        elif t["karne"] == "Az veri" and t["ret"] > 0:
            uy.append(uyari("Bilgi", "Az veride ret var", t["tedarikci"], f"{bilgi} Not vermek için en az {asgari} parti gerekir."))
    tedarikciler.sort(key=lambda t: (t["karne"] == "Az veri", -t["ppm"], katla(t["tedarikci"])))
    ay = {}
    for p in partiler:
        if p["tarih"]:
            ay.setdefault(f"{p['tarih']:%Y-%m}", []).append(p)
    return {"partiler": sorted(partiler, key=lambda p: (katla(p["tedarikci"]), p["tarih"] or date.min, p["parti"])), "tedarikciler": tedarikciler,
            "detay": sorted(kirilim("tedarikci", "malzeme"), key=lambda d: (katla(d["tedarikci"]), -d["ppm"], katla(d["malzeme"]))), "aylar": [{"ay": k, **gosterge(g)} for k, g in sorted(ay.items())], "toplam": gosterge(partiler),
            "notlar": {n: sum(1 for t in tedarikciler if t["karne"] == n) for n in NOT_RENK}, "uyarilar": uy}


BASLIK = ["Parti", "Gelen Miktar", "Ret Miktar", "Ret (ppm)", "Ret %", "Hatalı Parti", "Hatalı Parti %", "Reddedilen Parti", "Şartlı Kabul"]


def _deger(g: dict) -> list:
    return [g["parti"], g["gelen"], g["ret"], g["ppm"], g["ret_oran"], g["hatali"] or None, g["hatali_oran"], g["reddedilen"] or None, g["sartli"] or None]


def rapor_yaz(cikti: Path, r: dict, a: Decimal, b: Decimal, asgari: int) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Karne"
    _baslik(ws, ["Tedarikçi", "Not"] + BASLIK, (32, 9, 8, 15, 13, 12, 9, 10, 11, 11, 10), "B2")
    for t in r["tedarikciler"]:
        n = _satir(ws, [t["tedarikci"], t["karne"]] + _deger(t), para_s=(9,), dolgu={2: NOT_RENK[t["karne"]]})
        ws.cell(n, 6).number_format, ws.cell(n, 7).number_format = "#,##0", "0.000"
    n = _satir(ws, ["TOPLAM", None] + _deger(r["toplam"]), para_s=(9,))
    ws.cell(n, 6).number_format, ws.cell(n, 7).number_format = "#,##0", "0.000"
    for h in ws[n]:
        h.font = Font(bold=True)
    ws.append([])
    ws.append([f"Not eşikleri: A ≤ {a:,.0f} ppm, B ≤ {b:,.0f} ppm, C üstü; {asgari} partiden az veride not verilmez."])
    ws = wb.create_sheet("Tedarikçi × Malzeme")
    _baslik(ws, ["Tedarikçi", "Malzeme"] + BASLIK, (32, 26, 8, 15, 13, 12, 9, 10, 11, 11, 10), "C2")
    for d in r["detay"]:
        n = _satir(ws, [d["tedarikci"], d["malzeme"]] + _deger(d), para_s=(9,))
        ws.cell(n, 6).number_format, ws.cell(n, 7).number_format = "#,##0", "0.000"
    ws = wb.create_sheet("Aylar")
    _baslik(ws, ["Ay"] + BASLIK, (12, 8, 15, 13, 12, 9, 10, 11, 11, 10))
    for d in r["aylar"]:
        n = _satir(ws, [d["ay"]] + _deger(d), para_s=(8,))
        ws.cell(n, 5).number_format, ws.cell(n, 6).number_format = "#,##0", "0.000"
    ws = wb.create_sheet("Partiler")
    _baslik(ws, ["Tedarikçi", "Tarih", "Parti No", "Malzeme", "Gelen Miktar", "Ret Miktar", "Karar"], (32, 12, 16, 26, 14, 13, 14))
    for p in r["partiler"]:
        _satir(ws, [p["tedarikci"], p["tarih"], p["parti"], p["malzeme"], p["gelen"], p["ret"] or None, p["karar"]], tarih_s=(2,), dolgu={7: KIRMIZI if p["karar"] == "Ret" else SARI if p["karar"] == "Şartlı kabul" else None})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, a: Decimal = Decimal(1000), b: Decimal = Decimal(5000), asgari: int = 3) -> dict:
    partiler, uy = oku(yol)
    r = karne(partiler, a, b, asgari)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, a, b, asgari)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Giriş kalite kontrol verisinden tedarikçi bazında ret oranı (ppm) ve karne üretir.")
    p.add_argument("--kayitlar", type=Path, help="Tarih, Tedarikçi, Malzeme, Parti No, Gelen Miktar, Ret Miktar, Karar")
    p.add_argument("--a", default="1000", help="A notu için en yüksek ppm (varsayılan 1000)")
    p.add_argument("--b", default="5000", help="B notu için en yüksek ppm (varsayılan 5000)")
    p.add_argument("--min", type=int, default=3, dest="asgari", help="Not vermek için en az parti sayısı (varsayılan 3)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "tedarikci_kalite_karnesi.xlsx")
    args = p.parse_args(argv)
    yol = args.kayitlar or ORNEK / "giris_kontrol.csv"
    a, b = para(args.a), para(args.b)
    if a is None or b is None or not 0 <= a < b or args.asgari < 1:
        print("[X] --a ve --b ppm değeri olmalı ve a < b; --min en az 1 olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, args.cikti, a, b, args.asgari)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t, n = r["toplam"], r["notlar"]
    print(f"[OK] {len(r['tedarikciler'])} tedarikçi · {t['parti']} parti · {t['ppm']:,.0f} ppm · A {n['A']}, B {n['B']}, C {n['C']}, az veri {n['Az veri']}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {args.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
