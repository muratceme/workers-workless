"""
Kur Riski Pozisyon Raporu — Workers / Workless kod bloğu
Finans › Finans Müdürü

Döviz cinsinden varlık ve yükümlülüklerden net döviz pozisyonunu ve kur şoku senaryolarının etkisini hesaplar.
  Bilanço net pozisyonu = döviz varlıklar − döviz yükümlülükler (para birimi bazında)
  Türev net             = türev alım − türev satım (vadeli döviz alım/satım gibi riskten korunma işlemleri)
  Net pozisyon          = bilanço net pozisyonu + türev net;   artı = uzun (döviz fazlası), eksi = kısa (döviz açığı)
  TL karşılığı          = net pozisyon × kur (--kurlar dosyasından)
  Şok etkisi            = net pozisyon TL × şok %  (kur %10 artarsa kısa pozisyon zarar, uzun pozisyon kâr yazar)
  Pozisyon / özkaynak   = toplam net pozisyon TL ÷ --ozkaynak (verilirse); --limit verilirse mutlak oran bununla karşılaştırılır
  Vade kırılımı         = vadesi yazılan kalemler --tarih gününe göre 0–30, 31–90, 91–365 gün ve 1 yıldan uzun dilimlerine ayrılır
Hesap, girilen bakiyelerin anlık fotoğrafıdır; gelecekteki döviz gelir ve giderleri ile kurlar arası korelasyon hesaba katılmaz.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
SUTUN = {"kalem": ("kalem", "hesap", "aciklama"), "tur": ("tur", "tip", "sinif"), "pb": ("para birimi", "doviz", "doviz cinsi", "pb"), "tutar": ("tutar", "bakiye", "doviz tutari"), "vade": ("vade", "vade tarihi")}
KUR_SUTUN = {"pb": ("para birimi", "doviz", "doviz cinsi", "pb"), "kur": ("kur", "doviz kuru")}
TUR = {"varlik": "Varlık", "alacak": "Varlık", "aktif": "Varlık", "yukumluluk": "Yükümlülük", "borc": "Yükümlülük", "pasif": "Yükümlülük", "turev alim": "Türev alım", "vadeli alim": "Türev alım", "forward alim": "Türev alım",
       "turev satim": "Türev satım", "vadeli satim": "Türev satım", "forward satim": "Türev satım"}
ISARET = {"Varlık": 1, "Yükümlülük": -1, "Türev alım": 1, "Türev satım": -1}
DILIM = ("Vadesiz", "0–30 gün", "31–90 gün", "91–365 gün", "1 yıldan uzun")
Z, YUZ = Decimal(0), Decimal(100)


def kur_oku(yol: Path) -> dict[str, Decimal]:
    k = {}
    for r in kayitlar(yol, KUR_SUTUN, ("pb", "kur")):
        v = para(r["kur"])
        if metin(r["pb"]) and v and v > 0:
            k[metin(r["pb"]).upper()] = v
    return k


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("tur", "pb", "tutar")):
        ad, pb = metin(r["kalem"]), metin(r["pb"]).upper()
        if not pb and not metin(r["tutar"]):
            continue
        kim = f"{ad or '?'} (satır {r['_satir']})"
        t, tur = para(r["tutar"]), TUR.get(katla(r["tur"]))
        if not pb or t is None or t < 0 or tur is None:
            uy.append(uyari("Yüksek", "Satır okunamadı", kim, f"Tür '{metin(r['tur'])}', para birimi '{pb}', tutar '{metin(r['tutar'])}'. Tür: Varlık, Yükümlülük, Türev alım veya Türev satım; tutar artı işaretli olmalı. Satır rapora alınmadı."))
            continue
        if pb in ("TL", "TRY", "₺"):
            uy.append(uyari("Bilgi", "TL kalem atlandı", kim, "TL cinsinden kalem kur riski taşımaz; rapora alınmadı."))
            continue
        l.append({"kalem": ad, "tur": tur, "pb": pb, "tutar": t, "vade": tarih(r["vade"])})
    return l, uy


def dilim(vade: date | None, bugun: date) -> str:
    if vade is None:
        return DILIM[0]
    g = (vade - bugun).days
    return DILIM[1] if g <= 30 else DILIM[2] if g <= 90 else DILIM[3] if g <= 365 else DILIM[4]


def hesapla(kalemler: list[dict], kurlar: dict[str, Decimal], bugun: date, soklar: tuple = (Decimal(10), Decimal(20), Decimal(-10)), ozkaynak: Decimal | None = None, limit: Decimal | None = None) -> dict:
    if not kalemler:
        raise ValueError("Raporlanacak döviz kalemi yok.")
    uy, poz, vade = [], {}, {}
    for k in kalemler:
        p = poz.setdefault(k["pb"], {"pb": k["pb"], "Varlık": Z, "Yükümlülük": Z, "Türev alım": Z, "Türev satım": Z})
        p[k["tur"]] += k["tutar"]
        k["dilim"] = dilim(k["vade"], bugun)
        v = vade.setdefault(k["pb"], dict.fromkeys(DILIM, Z))
        v[k["dilim"]] += ISARET[k["tur"]] * k["tutar"]
        if k["vade"] and k["vade"] < bugun:
            uy.append(uyari("Bilgi", "Vadesi geçmiş kalem", f"{k['kalem']} · {k['pb']}", f"Vade {k['vade']:%d.%m.%Y} geçmiş; kalem 0–30 gün diliminde gösterildi. Bakiyenin güncel olduğunu doğrulayın."))
    pozisyonlar, net_tl, brut_tl = [], Z, Z
    for pb in sorted(poz):
        p = poz[pb]
        p["bilanco"], p["turev"] = p["Varlık"] - p["Yükümlülük"], p["Türev alım"] - p["Türev satım"]
        p["net"] = p["bilanco"] + p["turev"]
        p["yon"] = "Uzun (döviz fazlası)" if p["net"] > 0 else "Kısa (döviz açığı)" if p["net"] < 0 else "Dengede"
        p["kur"] = kurlar.get(pb)
        p["net_tl"] = yuvarla(p["net"] * p["kur"]) if p["kur"] else None
        p["soklar"] = {s: yuvarla(p["net_tl"] * s / YUZ) for s in soklar} if p["kur"] else {}
        if p["kur"] is None:
            uy.append(uyari("Yüksek", "Kur bulunamadı", pb, f"{pb} için kur verilmedi; {p['net']:,.2f} {pb} net pozisyon TL toplamına ve senaryolara girmedi."))
        else:
            net_tl, brut_tl = net_tl + p["net_tl"], brut_tl + abs(p["net_tl"])
            if p["turev"] and p["bilanco"] and (p["turev"] > 0) == (p["bilanco"] > 0):
                uy.append(uyari("Orta", "Türev pozisyonu riski büyütüyor", pb, f"Bilanço pozisyonu {p['bilanco']:,.2f}, türev net {p['turev']:,.2f} {pb}; ikisi aynı yönde. Korunma amaçlıysa yönünü kontrol edin."))
        pozisyonlar.append(p)
    oran = yuvarla(net_tl * YUZ / ozkaynak) if ozkaynak else None
    if oran is not None and limit is not None and abs(oran) > limit:
        uy.append(uyari("Yüksek", "Pozisyon limiti aşıldı", "Toplam", f"Net pozisyon / özkaynak %{oran}; belirlediğiniz sınır %{limit}."))
    return {"kalemler": kalemler, "pozisyonlar": pozisyonlar, "vade": {pb: vade[pb] for pb in sorted(vade)}, "net_tl": net_tl, "brut_tl": brut_tl, "soklar": {s: yuvarla(net_tl * s / YUZ) for s in soklar},
            "ozkaynak": ozkaynak, "oran": oran, "limit": limit, "bugun": bugun, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Pozisyon"
    _baslik(ws, ["Para Birimi", "Varlık", "Yükümlülük", "Bilanço Net", "Türev Alım", "Türev Satım", "Net Pozisyon", "Yön", "Kur", "Net Pozisyon (TL)"], (10, 17, 17, 17, 15, 15, 17, 22, 10, 19))
    for p in r["pozisyonlar"]:
        _satir(ws, [p["pb"], p["Varlık"], p["Yükümlülük"], p["bilanco"], p["Türev alım"] or None, p["Türev satım"] or None, p["net"], p["yon"], p["kur"], p["net_tl"]], para_s=(2, 3, 4, 5, 6, 7, 10),
               dolgu={8: KIRMIZI if p["net"] < 0 else YESIL if p["net"] > 0 else None})
    n = _satir(ws, ["TOPLAM (TL)", None, None, None, None, None, None, None, None, r["net_tl"]], para_s=(10,))
    ws.cell(n, 1).font = ws.cell(n, 10).font = Font(bold=True)
    ws.append([])
    if r["oran"] is not None:
        ws.append([f"Net pozisyon / özkaynak: %{r['oran']} (özkaynak {r['ozkaynak']:,.2f} TL)" + (f" · sınır %{r['limit']}" if r["limit"] is not None else "")])
    ws.append([f"Durum tarihi {r['bugun']:%d.%m.%Y}. Artı = döviz fazlası (kur artışında kâr), eksi = döviz açığı (kur artışında zarar)."])
    ws = wb.create_sheet("Senaryolar")
    soklar = list(r["soklar"])
    _baslik(ws, ["Para Birimi", "Net Pozisyon (TL)"] + [f"Kur %{s:+} olursa etki (TL)" for s in soklar], [14, 19] + [22] * len(soklar))
    for p in r["pozisyonlar"]:
        if p["kur"]:
            _satir(ws, [p["pb"], p["net_tl"]] + [p["soklar"][s] for s in soklar], para_s=range(2, 3 + len(soklar)), dolgu={3 + i: KIRMIZI for i, s in enumerate(soklar) if p["soklar"][s] < 0})
    n = _satir(ws, ["TOPLAM", r["net_tl"]] + list(r["soklar"].values()), para_s=range(2, 3 + len(soklar)), dolgu={3 + i: KIRMIZI for i, s in enumerate(soklar) if r["soklar"][s] < 0})
    for h in ws[n]:
        h.font = Font(bold=True)
    ws.append([])
    ws.append(["Tüm kurların TL karşısında aynı oranda değiştiği varsayılır. Etki, kur farkı kârı (+) veya zararıdır (−); vergi ve korunma muhasebesi etkileri dahil değildir."])
    ws = wb.create_sheet("Vade Kırılımı")
    _baslik(ws, ["Para Birimi"] + list(DILIM) + ["Net"], [12] + [17] * 6)
    for pb, v in r["vade"].items():
        _satir(ws, [pb] + [v[d] or None for d in DILIM] + [sum(v.values(), Z)], para_s=range(2, 8), dolgu={2 + i: KIRMIZI for i, d in enumerate(DILIM) if v[d] < 0})
    ws = wb.create_sheet("Kalemler")
    _baslik(ws, ["Kalem", "Tür", "Para Birimi", "Tutar", "Vade", "Vade Dilimi"], (36, 14, 10, 17, 12, 16))
    for k in r["kalemler"]:
        _satir(ws, [k["kalem"], k["tur"], k["pb"], k["tutar"], k["vade"], k["dilim"]], para_s=(4,), tarih_s=(5,))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, kur_yol: Path, cikti: Path, bugun: date, soklar: tuple = (Decimal(10), Decimal(20), Decimal(-10)), ozkaynak: Decimal | None = None, limit: Decimal | None = None) -> dict:
    kalemler, uy = oku(yol)
    r = hesapla(kalemler, kur_oku(kur_yol), bugun, soklar, ozkaynak, limit)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Net döviz pozisyonunu ve kur şoku senaryolarını hesaplar.")
    p.add_argument("--kalemler", type=Path, help="Kalem, Tür (Varlık / Yükümlülük / Türev alım / Türev satım), Para Birimi, Tutar, Vade")
    p.add_argument("--kurlar", type=Path, help="Para Birimi, Kur")
    p.add_argument("--sok", default="10;20;-10", help="Kur şoku senaryoları, %% (noktalı virgülle; varsayılan 10;20;-10)")
    p.add_argument("--ozkaynak", help="Özkaynak tutarı (TL); verilirse net pozisyon / özkaynak oranı hesaplanır")
    p.add_argument("--limit", help="Net pozisyon / özkaynak için kendi belirlediğiniz sınır, %% (verilirse aşım uyarılır)")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kur_riski_pozisyon.xlsx")
    a = p.parse_args(argv)
    ornek = a.kalemler is None
    yol, kur_yol = (ORNEK / "kalemler.csv", a.kurlar or ORNEK / "kurlar.csv") if ornek else (a.kalemler, a.kurlar)
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    soklar = tuple(para(x) for x in re.split(r"[;\s]+", a.sok.strip()) if x)
    ozkaynak, limit = para(a.ozkaynak) if a.ozkaynak else Decimal(60000000) if ornek else None, para(a.limit) if a.limit else None
    if bugun is None or not soklar or any(s is None or s == 0 or s <= -100 for s in soklar) or len(set(soklar)) != len(soklar) or (a.ozkaynak and (ozkaynak is None or ozkaynak <= 0)) or (a.limit and (limit is None or limit <= 0)):
        print("[X] --tarih GG.AA.YYYY; --sok noktalı virgülle ayrılmış, sıfırdan farklı yüzdeler (örn. --sok=10;20;-10); --ozkaynak ve --limit pozitif sayı olmalı")
        return 2
    if kur_yol is None:
        print("[X] --kurlar zorunludur (Para Birimi, Kur)")
        return 2
    for y in (yol, kur_yol):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(yol, kur_yol, a.cikti, bugun, soklar, ozkaynak, limit)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print("[OK] " + " · ".join(f"{x['pb']} net {x['net']:,.2f}" for x in r["pozisyonlar"]) + f" · toplam {r['net_tl']:,.2f} TL" + (f" · özkaynağın %{r['oran']}" if r["oran"] is not None else ""))
    print("[OK] " + " · ".join(f"kur %{s:+} → {e:,.2f} TL" for s, e in r["soklar"].items()))
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
