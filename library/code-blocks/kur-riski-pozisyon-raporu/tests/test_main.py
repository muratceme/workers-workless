import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)
KUR = {"USD": D(40), "EUR": D(50)}


def k(tur, pb, tutar, vade=None, ad="K"):
    return {"kalem": ad, "tur": tur, "pb": pb, "tutar": D(tutar), "vade": vade}


class HesapTesti(unittest.TestCase):
    def test_pozisyon_ve_sok(self):
        l = [k("Varlık", "USD", 1000), k("Yükümlülük", "USD", 400), k("Varlık", "EUR", 100), k("Yükümlülük", "EUR", 900), k("Türev alım", "EUR", 300), k("Türev satım", "EUR", 100), k("Varlık", "GBP", 50), k("Yükümlülük", "GBP", 50)]
        r = main.hesapla(l, KUR | {"GBP": D(55)}, BUGUN, (D(10), D(-25)), D(100000))
        self.assertEqual([(p["pb"], p["bilanco"], p["turev"], p["net"], p["yon"], p["net_tl"], p["soklar"]) for p in r["pozisyonlar"]],
                         [("EUR", D(-800), D(200), D(-600), "Kısa (döviz açığı)", D("-30000.00"), {D(10): D("-3000.00"), D(-25): D("7500.00")}),
                          ("GBP", D(0), D(0), D(0), "Dengede", D("0.00"), {D(10): D("0.00"), D(-25): D("0.00")}),
                          ("USD", D(600), D(0), D(600), "Uzun (döviz fazlası)", D("24000.00"), {D(10): D("2400.00"), D(-25): D("-6000.00")})])
        self.assertEqual((r["net_tl"], r["brut_tl"], r["soklar"], r["oran"], r["uyarilar"]), (D(-6000), D(54000), {D(10): D("-600.00"), D(-25): D("1500.00")}, D("-6.00"), []))

    def test_uyarilar_ve_limit(self):
        l = [k("Yükümlülük", "EUR", 1000), k("Türev satım", "EUR", 200), k("Varlık", "CHF", 10)]
        r = main.hesapla(l, KUR, BUGUN, (D(10),), D(100000), D(50))
        self.assertEqual((r["pozisyonlar"][0]["net_tl"], r["pozisyonlar"][1]["net"], r["net_tl"], r["oran"], sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"])),
                         (None, D(-1200), D(-60000), D("-60.00"), [("Orta", "Türev pozisyonu riski büyütüyor", "EUR"), ("Yüksek", "Kur bulunamadı", "CHF"), ("Yüksek", "Pozisyon limiti aşıldı", "Toplam")]))
        self.assertEqual([u["tur"] for u in main.hesapla(l[:1], KUR, BUGUN, (D(10),), D(100000), D(50))["uyarilar"]], [])
        self.assertEqual(main.hesapla(l[:1], KUR, BUGUN)["oran"], None)
        with self.assertRaises(ValueError):
            main.hesapla([], KUR, BUGUN)

    def test_vade_dilimi(self):
        g = lambda n: BUGUN + timedelta(days=n)   # noqa: E731
        self.assertEqual([main.dilim(g(n), BUGUN) for n in (-5, 0, 30, 31, 90, 91, 365, 366)] + [main.dilim(None, BUGUN)],
                         ["0–30 gün", "0–30 gün", "0–30 gün", "31–90 gün", "31–90 gün", "91–365 gün", "91–365 gün", "1 yıldan uzun", "Vadesiz"])
        r = main.hesapla([k("Varlık", "USD", 100), k("Varlık", "USD", 50, g(10)), k("Yükümlülük", "USD", 70, g(20)), k("Türev satım", "USD", 30, g(100)), k("Yükümlülük", "USD", 5, g(-3), "Eski")], KUR, BUGUN)
        self.assertEqual((r["vade"], [(u["tur"], u["kim"]) for u in r["uyarilar"]]),
                         ({"USD": {"Vadesiz": D(100), "0–30 gün": D(-25), "31–90 gün": D(0), "91–365 gün": D(-30), "1 yıldan uzun": D(0)}}, [("Vadesi geçmiş kalem", "Eski · USD")]))


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "k.xlsx"
            r = main.calistir(ORNEK / "kalemler.csv", ORNEK / "kurlar.csv", cikti, BUGUN, ozkaynak=D(60000000))
            self.assertEqual([(p["pb"], p["Varlık"], p["Yükümlülük"], p["bilanco"], p["turev"], p["net"], p["net_tl"]) for p in r["pozisyonlar"]],
                             [("EUR", D(400000), D(1650000), D(-1250000), D(500000), D(-750000), D(-36750000)), ("USD", D(850000), D(650000), D(200000), D(0), D(200000), D(8400000))])
            self.assertEqual((r["net_tl"], r["soklar"], r["oran"], r["pozisyonlar"][0]["soklar"][D(10)], r["uyarilar"]),
                             (D(-28350000), {D(10): D(-2835000), D(20): D(-5670000), D(-10): D(2835000)}, D("-47.25"), D(-3675000), []))
            self.assertEqual(r["vade"], {"EUR": {"Vadesiz": D(0), "0–30 gün": D(0), "31–90 gün": D(100000), "91–365 gün": D(-700000), "1 yıldan uzun": D(-150000)},
                                         "USD": {"Vadesiz": D(250000), "0–30 gün": D(-350000), "31–90 gün": D(600000), "91–365 gün": D(-300000), "1 yıldan uzun": D(0)}})
            self.assertEqual(load_workbook(cikti).sheetnames, ["Pozisyon", "Senaryolar", "Vade Kırılımı", "Kalemler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, kr = Path(tmp) / "k.csv", Path(tmp) / "kur.csv"
            y.write_text("Hesap;Tip;Döviz;Bakiye;Vade Tarihi\nMevduat;AKTİF;usd;1.250,50;\nKredi;Borç;eur;2.000;30.06.2027\nForward;Vadeli Satım;usd;500;x\nKasa;Varlık;TL;100;\nX;Stok;USD;5;\nY;Varlık;USD;-5;\nZ;Varlık;;5;\n",
                         encoding="utf-8")
            kr.write_text("Döviz;Kur\nusd;40\neur;0\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["kalem"], x["tur"], x["pb"], x["tutar"], x["vade"]) for x in l], [u["tur"] for u in uy]),
                             ([("Mevduat", "Varlık", "USD", D("1250.50"), None), ("Kredi", "Yükümlülük", "EUR", D(2000), date(2027, 6, 30)), ("Forward", "Türev satım", "USD", D(500), None)],
                              ["TL kalem atlandı"] + ["Satır okunamadı"] * 3))
            self.assertEqual(main.kur_oku(kr), {"USD": D(40)})
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--kalemler", str(y), "--kurlar", str(kr), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "12.10.2026", "--ozkaynak", "1.000.000", "--limit", "30", "--sok=-5;5;12.5"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + x) for x in (["--sok", "0"], ["--sok", "x"], ["--sok", "10;10"], ["--sok=-100"], ["--ozkaynak", "0"], ["--limit", "x"], ["--tarih", "x"])], [2] * 7)
            self.assertEqual((main.main(["--kalemler", str(y)]), main.main(["--kalemler", str(Path(tmp) / "yok.csv"), "--kurlar", str(kr)]), main.main(["--kalemler", str(y), "--kurlar", str(Path(tmp) / "yok.csv")])), (2, 1, 1))
            for icerik in ("Kalem;Tutar\nA;1\n", "Tür;Para Birimi;Tutar\nx;USD;1\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
