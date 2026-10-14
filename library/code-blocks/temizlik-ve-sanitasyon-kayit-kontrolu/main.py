"""
Temizlik ve Sanitasyon Kayıt Kontrolü — Workers / Workless kod bloğu
Üretim › Üretim Şefi
Hat temizlik ve sanitasyon kayıtlarını plana göre kontrol edip eksik, geciken ve uygunsuz kayıtları listeler.
  Sıklık        = Vardiya (günde --vardiya kez) · Günlük · Haftalık (Pzt-Paz) · Aylık · "N gün" (iki temizlik arası en çok N gün)
  Beklenen      = dönemdeki çalışma günü / hafta / ay sayısına göre; --kapali günler (ör. "Paz") vardiya ve günlük görevlerde sayılmaz
  Uyum %        = kaydı bulunan periyot ÷ beklenen periyot × 100   (vardiyada gün başına en çok --vardiya kayıt sayılır)
  Kontroller    = eksik gün / hafta / ay · N gün aralığının aşılması · sonucu uygunsuz olup düzeltici faaliyeti yazılmamış kayıt ·
                  planda olmayan görev · onaylanmamış kayıt
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

PLAN = {"hat": ("hat", "alan", "bolum", "hat alan"), "gorev": ("gorev", "temizlik gorevi", "ekipman", "islem"), "siklik": ("siklik", "periyot", "temizlik sikligi")}
KAYIT = {"tarih": ("tarih", "temizlik tarihi"), "vardiya": ("vardiya",), "hat": ("hat", "alan", "bolum", "hat alan"), "gorev": ("gorev", "temizlik gorevi", "ekipman", "islem"),
         "yapan": ("yapan", "uygulayan", "personel"), "onay": ("kontrol eden", "onaylayan", "onay"), "sonuc": ("sonuc", "durum", "dogrulama", "atp sonucu"),
         "dof": ("duzeltici faaliyet", "duzeltici islem", "aciklama", "dof")}
GUNLER = {"pzt": 0, "pazartesi": 0, "sal": 1, "sali": 1, "car": 2, "carsamba": 2, "per": 3, "persembe": 3, "cum": 4, "cuma": 4, "cmt": 5, "cumartesi": 5, "paz": 6, "pazar": 6}


def siklik_oku(x) -> tuple[str, int] | None:
    k = katla(x)
    if not k:
        return None
    if "vardiya" in k:
        return ("Vardiya", 1)
    if k in ("gunluk", "her gun", "gun"):
        return ("Günlük", 1)
    if "haftalik" in k or k == "her hafta":
        return ("Haftalık", 7)
    if "aylik" in k or k == "her ay":
        return ("Aylık", 30)
    m = re.fullmatch(r"(?:her )?(\d+) gun(?:de bir)?", k)
    return ("N gün", int(m.group(1))) if m and int(m.group(1)) > 0 else None


def kapali_oku(x: str) -> set[int] | None:
    s = set()
    for p in re.split(r"[;,\s]+", x.strip()):
        if not p:
            continue
        g = GUNLER.get(katla(p))
        if g is None:
            return None
        s.add(g)
    return s


def oku_plan(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, PLAN, ("gorev", "siklik")):
        g = metin(r["gorev"])
        if not g:
            continue
        s = siklik_oku(r["siklik"])
        if s is None:
            uy.append(uyari("Yüksek", "Plan satırı okunamadı", g, f"Satır {r['_satir']}: sıklık '{metin(r['siklik'])}' anlaşılamadı (Vardiya, Günlük, Haftalık, Aylık ya da '3 gün')."))
            continue
        l.append({"hat": metin(r["hat"]) or "-", "gorev": g, "siklik": s[0], "n": s[1]})
    return l, uy


def oku_kayit(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, KAYIT, ("tarih", "gorev")):
        g = metin(r["gorev"])
        if not g:
            continue
        t = tarih(r["tarih"])
        if t is None:
            uy.append(uyari("Yüksek", "Kayıt satırı okunamadı", g, f"Satır {r['_satir']}: tarih okunamadı. Kayıt alınmadı."))
            continue
        sk = katla(r["sonuc"])
        sonuc = "Uygun değil" if ("degil" in sk or sk in ("red", "ret", "basarisiz", "uygunsuz", "kaldi")) else "Uygun" if sk else ""
        l.append({"tarih": t, "vardiya": metin(r["vardiya"]), "hat": metin(r["hat"]) or "-", "gorev": g, "yapan": metin(r["yapan"]), "onay": metin(r["onay"]), "sonuc": sonuc, "dof": metin(r["dof"])})
    return l, uy


def kontrol(plan: list[dict], kayit: list[dict], bas: date, bit: date, vardiya: int = 1, kapali: set[int] | None = None) -> dict:
    if not plan:
        raise ValueError("Temizlik planı boş.")
    if bit < bas:
        raise ValueError("Dönem sonu dönem başından önce olamaz.")
    kapali = kapali or set()
    uy, sonuc = [], []
    gunler = [bas + timedelta(days=i) for i in range((bit - bas).days + 1)]
    calisma = [g for g in gunler if g.weekday() not in kapali]
    donem = [k for k in kayit if bas <= k["tarih"] <= bit]
    planli = set()
    for p in plan:
        anahtar = (katla(p["hat"]), katla(p["gorev"]))
        planli.add(anahtar)
        ks = sorted((k for k in donem if (katla(k["hat"]), katla(k["gorev"])) == anahtar or (p["hat"] == "-" and katla(k["gorev"]) == anahtar[1])), key=lambda k: k["tarih"])
        say = defaultdict(int)
        for k in ks:
            say[k["tarih"]] += 1
        eksik = []
        if p["siklik"] in ("Vardiya", "Günlük"):
            hedef = vardiya if p["siklik"] == "Vardiya" else 1
            beklenen = len(calisma) * hedef
            yapilan = sum(min(say[g], hedef) for g in calisma)
            eksik = [f"{g:%d.%m}" + (f" ({hedef - say[g]})" if hedef > 1 else "") for g in calisma if say[g] < hedef]
        elif p["siklik"] in ("Haftalık", "Aylık"):
            fn = (lambda g: g.isocalendar()[:2]) if p["siklik"] == "Haftalık" else (lambda g: (g.year, g.month))
            periyot = sorted({fn(g) for g in gunler})
            var = {fn(k["tarih"]) for k in ks}
            beklenen, yapilan = len(periyot), sum(1 for x in periyot if x in var)
            eksik = [(f"{x[0]}-H{x[1]:02d}" if p["siklik"] == "Haftalık" else f"{x[1]:02d}.{x[0]}") for x in periyot if x not in var]
        else:
            beklenen = math.ceil(len(gunler) / p["n"])
            yapilan = min(len({k["tarih"] for k in ks}), beklenen)
            onceki = bas - timedelta(days=1)
            for g in sorted({k["tarih"] for k in ks}) + [bit + timedelta(days=1)]:
                if (g - onceki).days > p["n"]:
                    eksik.append(f"{onceki + timedelta(days=1):%d.%m}–{g - timedelta(days=1):%d.%m} ({(g - onceki).days - 1} gün kayıtsız)" if g <= bit else f"{onceki:%d.%m} sonrası kayıt yok")
                onceki = g
        uyum = yuvarla(Decimal(yapilan) * 100 / beklenen) if beklenen else None
        s = {**p, "beklenen": beklenen, "yapilan": yapilan, "uyum": uyum, "eksik": eksik, "kayit": len(ks), "uygunsuz": sum(1 for k in ks if k["sonuc"] == "Uygun değil")}
        sonuc.append(s)
        if eksik:
            uy.append(uyari("Orta", "Eksik ya da geciken temizlik kaydı", f"{p['hat']} · {p['gorev']}", f"{p['siklik']}" + (f" ({p['n']} gün)" if p["siklik"] == "N gün" else "")
                            + f": uyum %{uyum}; " + ", ".join(eksik[:10]) + (" …" if len(eksik) > 10 else "") + "."))
    for k in donem:
        kim = f"{k['hat']} · {k['gorev']} · {k['tarih']:%d.%m.%Y}"
        if k["sonuc"] == "Uygun değil" and not k["dof"]:
            uy.append(uyari("Yüksek", "Uygunsuz sonuçta düzeltici faaliyet yok", kim, "Doğrulama sonucu uygun değil; yeniden temizlik ve düzeltici faaliyet kaydı bulunamadı."))
    genel = {g for h, g in planli if not h}
    plansiz = sorted({(k["hat"], k["gorev"]) for k in donem if (katla(k["hat"]), katla(k["gorev"])) not in planli and katla(k["gorev"]) not in genel}, key=lambda x: (katla(x[0]), katla(x[1])))
    for h, g in plansiz:
        uy.append(uyari("Bilgi", "Planda olmayan görev", f"{h} · {g}", "Kayıt var ama temizlik planında yok; plana ekleyin ya da yazımı düzeltin."))
    onaysiz = sum(1 for k in donem if not k["onay"])
    if onaysiz:
        uy.append(uyari("Bilgi", "Onaylanmamış kayıt", f"{onaysiz} kayıt", "Kontrol eden / onaylayan alanı boş."))
    sonuc.sort(key=lambda s: (s["uyum"] if s["uyum"] is not None else Decimal(101), katla(s["hat"]), katla(s["gorev"])))
    toplam_b, toplam_y = sum(s["beklenen"] for s in sonuc), sum(s["yapilan"] for s in sonuc)
    return {"gorevler": sonuc, "kayitlar": sorted(donem, key=lambda k: (k["tarih"], katla(k["hat"]), katla(k["gorev"]))), "uyum": yuvarla(Decimal(toplam_y) * 100 / toplam_b) if toplam_b else None,
            "calisma_gunu": len(calisma), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, bas: date, bit: date) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Görev Uyumu"
    _baslik(ws, ["Hat / Alan", "Görev", "Sıklık", "Beklenen", "Yapılan", "Uyum %", "Kayıt", "Uygunsuz", "Eksik / Geciken"], (16, 28, 10, 9, 9, 8, 7, 9, 50), "C2")
    for s in r["gorevler"]:
        _satir(ws, [s["hat"], s["gorev"], s["siklik"] + (f" ({s['n']})" if s["siklik"] == "N gün" else ""), s["beklenen"], s["yapilan"], s["uyum"], s["kayit"], s["uygunsuz"] or None, ", ".join(s["eksik"])],
               para_s=(6,), dolgu={6: None if s["uyum"] is None else YESIL if s["uyum"] >= 100 else SARI if s["uyum"] >= 90 else KIRMIZI, 8: KIRMIZI if s["uygunsuz"] else None}, sar=(9,))
    ws.append([])
    ws.append([f"Dönem {bas:%d.%m.%Y} – {bit:%d.%m.%Y}; çalışma günü {r['calisma_gunu']}. Genel uyum %{r['uyum']}."])
    ws = wb.create_sheet("Kayıtlar")
    _baslik(ws, ["Tarih", "Vardiya", "Hat / Alan", "Görev", "Yapan", "Kontrol Eden", "Sonuç", "Düzeltici Faaliyet"], (11, 8, 16, 28, 16, 16, 11, 34), "D2")
    for k in r["kayitlar"]:
        _satir(ws, [k["tarih"], k["vardiya"], k["hat"], k["gorev"], k["yapan"], k["onay"], k["sonuc"], k["dof"]], tarih_s=(1,),
               dolgu={7: KIRMIZI if k["sonuc"] == "Uygun değil" else None, 6: SARI if not k["onay"] else None}, sar=(8,))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(plan_yol: Path, kayit_yol: Path, cikti: Path, bas: date, bit: date, vardiya: int = 1, kapali: set[int] | None = None) -> dict:
    plan, uy1 = oku_plan(plan_yol)
    kayit, uy2 = oku_kayit(kayit_yol)
    r = kontrol(plan, kayit, bas, bit, vardiya, kapali)
    r["uyarilar"] = sirala(uy1 + uy2 + r["uyarilar"])
    rapor_yaz(cikti, r, bas, bit)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Hat temizlik ve sanitasyon kayıtlarını plana göre kontrol edip eksik ya da geciken kayıtları listeler.")
    p.add_argument("--plan", type=Path, help="Hat / Alan, Görev, Sıklık (Vardiya, Günlük, Haftalık, Aylık, '3 gün')")
    p.add_argument("--kayitlar", type=Path, help="Tarih, Vardiya, Hat / Alan, Görev, Yapan, Kontrol Eden, Sonuç, Düzeltici Faaliyet")
    p.add_argument("--baslangic", help="Dönem başı GG.AA.YYYY (varsayılan kayıtların ilk günü)")
    p.add_argument("--bitis", help="Dönem sonu GG.AA.YYYY (varsayılan kayıtların son günü)")
    p.add_argument("--vardiya", type=int, default=1, help="'Vardiya' sıklıklı görevlerde günlük vardiya sayısı (varsayılan 1)")
    p.add_argument("--kapali", default="", help="Üretim yapılmayan günler, ör. \"Paz\" ya da \"Cmt,Paz\"")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "sanitasyon_kayit_kontrolu.xlsx")
    a = p.parse_args(argv)
    ornek = a.plan is None
    py, ky = a.plan or ORNEK / "temizlik_plani.csv", a.kayitlar or ORNEK / "temizlik_kayitlari.csv"
    if ornek and a.vardiya == 1 and not a.kapali:
        a.vardiya, a.kapali = 2, "Paz"
    kapali = kapali_oku(a.kapali)
    if a.vardiya < 1 or kapali is None:
        print("[X] --vardiya en az 1; --kapali gün adları (Pzt, Sal, Çar, Per, Cum, Cmt, Paz) olmalı")
        return 2
    for y in (py, ky):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        plan, uy1 = oku_plan(py)
        kayit, uy2 = oku_kayit(ky)
        bas = tarih(a.baslangic) if a.baslangic else min((k["tarih"] for k in kayit), default=None)
        bit = tarih(a.bitis) if a.bitis else max((k["tarih"] for k in kayit), default=None)
        if bas is None or bit is None:
            print("[X] Dönem belirlenemedi: --baslangic / --bitis GG.AA.YYYY olmalı ya da kayıt dosyasında tarih bulunmalı")
            return 2
        r = kontrol(plan, kayit, bas, bit, a.vardiya, kapali)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    r["uyarilar"] = sirala(uy1 + uy2 + r["uyarilar"])
    rapor_yaz(a.cikti, r, bas, bit)
    print(f"[OK] {len(r['gorevler'])} görev · {len(r['kayitlar'])} kayıt · {bas:%d.%m.%Y}–{bit:%d.%m.%Y} · genel uyum %{r['uyum']}")
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
