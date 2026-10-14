import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BAS, BIT = date(2026, 10, 5), date(2026, 10, 11)


def pl(hat, gorev, siklik, n=1):
    return {"hat": hat, "gorev": gorev, "siklik": siklik, "n": n}


def k(gun, hat, gorev, sonuc="Uygun", onay="Ş", dof=""):
    return {"tarih": date(2026, 10, gun), "vardiya": "", "hat": hat, "gorev": gorev, "yapan": "", "onay": onay, "sonuc": sonuc, "dof": dof}


PLAN = [pl("H1", "Konveyör yıkama", "Günlük"), pl("H1", "Zemin dezenfeksiyon", "Vardiya"), pl("H2", "Soğuk oda", "Haftalık"), pl("H2", "Drenaj", "N gün", 3)]


def kayit():
    l = [k(g, "H1", "Konveyör yıkama", "Uygun değil" if g == 9 else "Uygun") for g in (5, 6, 7, 9, 10)]
    l += [k(g, "H1", "Zemin dezenfeksiyon") for g in (5, 5, 6, 7, 7, 8, 8, 9, 9, 10, 10)]
    l += [k(7, "H2", "Soğuk oda"), k(5, "H2", "Drenaj"), k(9, "h2", "drenaj", onay=""), k(6, "H1", "Tavan temizliği"), k(20, "H1", "Konveyör yıkama")]
    return l


class KontrolTesti(unittest.TestCase):
    def test_uyum(self):
        r = main.kontrol(PLAN, kayit(), BAS, BIT, 2, {6})
        self.assertEqual([(s["gorev"], s["beklenen"], s["yapilan"], s["uyum"], s["eksik"]) for s in r["gorevler"]],
                         [("Drenaj", 3, 2, D("66.67"), ["06.10–08.10 (3 gün kayıtsız)"]), ("Konveyör yıkama", 6, 5, D("83.33"), ["08.10"]),
                          ("Zemin dezenfeksiyon", 12, 11, D("91.67"), ["06.10 (1)"]), ("Soğuk oda", 1, 1, D("100.00"), [])])
        self.assertEqual((r["uyum"], r["calisma_gunu"], len(r["kayitlar"])), (D("86.36"), 6, 20))
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Bilgi", "Onaylanmamış kayıt", "1 kayıt"), ("Bilgi", "Planda olmayan görev", "H1 · Tavan temizliği"), ("Orta", "Eksik ya da geciken temizlik kaydı", "H1 · Konveyör yıkama"),
                          ("Orta", "Eksik ya da geciken temizlik kaydı", "H1 · Zemin dezenfeksiyon"), ("Orta", "Eksik ya da geciken temizlik kaydı", "H2 · Drenaj"),
                          ("Yüksek", "Uygunsuz sonuçta düzeltici faaliyet yok", "H1 · Konveyör yıkama · 09.10.2026")])

    def test_yardimcilar(self):
        self.assertEqual([main.siklik_oku(x) for x in ("Her vardiya", "Günlük", "haftalık", "Aylık", "3 gün", "her 10 günde bir", "0 gün", "")],
                         [("Vardiya", 1), ("Günlük", 1), ("Haftalık", 7), ("Aylık", 30), ("N gün", 3), ("N gün", 10), None, None])
        self.assertEqual([main.kapali_oku(x) for x in ("", "Paz", "Cmt, Pazar", "x")], [set(), {6}, {5, 6}, None])
        r = main.kontrol([pl("-", "Aylık bakım", "Aylık")], [k(5, "H9", "aylık bakım")], date(2026, 9, 20), BIT)
        self.assertEqual(([(s["beklenen"], s["yapilan"], s["eksik"]) for s in r["gorevler"]], [u["tur"] for u in r["uyarilar"]]), ([(2, 1, ["09.2026"])], ["Eksik ya da geciken temizlik kaydı"]))
        for a, b in (([], BIT), (PLAN, date(2026, 10, 1))):
            with self.assertRaises(ValueError):
                main.kontrol(a, [], BAS, b)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "temizlik_plani.csv", ORNEK / "temizlik_kayitlari.csv", cikti, date(2026, 9, 28), date(2026, 10, 11), 2, {6})
            self.assertEqual(([(s["gorev"], s["uyum"]) for s in r["gorevler"]], r["uyum"]),
                             ([("Evaporatör temizliği", D("50.00")), ("Soğuk oda genel temizlik", D("50.00")), ("Drenaj kanalları", D("80.00")), ("Tartım alanı temizliği", D("83.33")),
                               ("Konveyör yıkama ve dezenfeksiyon", D("91.67")), ("Dolum nozulları CIP", D("95.83")), ("Zemin dezenfeksiyonu", D("95.83")), ("Karıştırma tankı CIP", D("100.00"))], D("91.40")))
            self.assertEqual([u["tur"] for u in r["uyarilar"]].count("Eksik ya da geciken temizlik kaydı"), 7)
            self.assertEqual([(u["tur"], u["kim"]) for u in r["uyarilar"] if u["onem"] != "Orta"],
                             [("Uygunsuz sonuçta düzeltici faaliyet yok", "Dolum hattı 1 · Dolum nozulları CIP · 07.10.2026"), ("Planda olmayan görev", "Dolum hattı 1 · Dolum makinesi dış yüzey")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Görev Uyumu", "Kayıtlar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, z = Path(tmp) / "p.csv", Path(tmp) / "k.csv"
            y.write_text("Alan;Görev;Periyot\nA;Yıkama;Günlük\nA;Silme;bazen\n;Yer;Haftalık\n", encoding="utf-8")
            z.write_text("Tarih;Alan;Görev;Sonuç;Onaylayan\n05.10.2026;A;Yıkama;Uygun değil;X\nx;A;Yıkama;;\n06.10.2026;A;Yıkama;;\n", encoding="utf-8")
            l, uy = main.oku_plan(y)
            self.assertEqual(([(x["hat"], x["gorev"], x["siklik"]) for x in l], len(uy)), ([("A", "Yıkama", "Günlük"), ("-", "Yer", "Haftalık")], 1))
            l, uy = main.oku_kayit(z)
            self.assertEqual(([(x["tarih"], x["sonuc"]) for x in l], len(uy)), ([(date(2026, 10, 5), "Uygun değil"), (date(2026, 10, 6), "")], 1))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--plan", str(y), "--kayitlar", str(z), "--cikti", c]
            self.assertEqual(main.main(arg + ["--baslangic", "05.10.2026", "--bitis", "11.10.2026", "--vardiya", "3", "--kapali", "Cmt,Paz"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--vardiya", "0"], ["--kapali", "x"], ["--baslangic", "x"])], [2] * 3)
            self.assertEqual(main.main(arg + ["--baslangic", "11.10.2026", "--bitis", "05.10.2026"]), 1)
            self.assertEqual(main.main(arg[:2] + ["--kayitlar", str(Path(tmp) / "yok.csv")]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
