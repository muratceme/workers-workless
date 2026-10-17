"""
Kampanya Fiyat Simülasyonu — Gerisi Bende kod bloğu
Pazaryeri Yönetimi › Pazaryeri Uzmanı
Planlanan indirim ve kupon kurgularının ürün bazında kâr marjına etkisini hesaplar (mevcut fiyat ↔ kampanya fiyatı).
  Müşterinin ödediği (KDV dahil) = satış fiyatı × (1 − indirim %) − kupon
  Net satış (KDV hariç)          = müşterinin ödediği ÷ (1 + KDV %)
  Komisyon                       = komisyon % × müşterinin ödediği (--matrah dahil) ya da × net satış (--matrah haric)
  Birim kâr                      = net satış − komisyon − kargo − hizmet bedeli − birim maliyet   (gider kalemleri KDV hariç)
  Marj %                         = birim kâr ÷ net satış × 100
  Gerekli adet artışı %          = (mevcut birim kâr ÷ kampanya birim kârı − 1) × 100 — toplam kârı korumak için
  Başabaş fiyat                  = kârın sıfır olduğu müşteri fiyatı; Min. marj fiyatı = marjın --min-marj (%10) olduğu fiyat
Kupon satıcı tarafından karşılanıyorsa yazılır; pazaryerinin karşıladığı kuponu yazmayın.
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"kod": ("urun kodu", "stok kodu", "barkod", "kod", "sku"), "ad": ("urun adi", "urun", "aciklama"), "fiyat": ("satis fiyati", "fiyat", "kdv dahil fiyat", "liste fiyati"),
         "kdv": ("kdv", "kdv orani", "kdv yuzdesi"), "maliyet": ("birim maliyet", "maliyet", "alis fiyati"), "komisyon": ("komisyon", "komisyon orani", "komisyon yuzdesi"),
         "kargo": ("kargo", "kargo bedeli", "kargo maliyeti"), "hizmet": ("hizmet bedeli", "islem bedeli", "diger gider"), "adet": ("mevcut adet", "aylik adet", "satis adedi", "adet"),
         "indirim": ("indirim", "indirim orani", "indirim yuzdesi", "kampanya indirimi"), "kupon": ("kupon", "kupon tutari", "sepet indirimi"), "beklenen": ("beklenen adet", "kampanya adedi", "hedef adet")}
SAYI = ("fiyat", "kdv", "maliyet", "komisyon", "kargo", "hizmet", "adet", "indirim", "kupon", "beklenen")
SIFIR = Decimal(0)


def oku(yol: Path, kdv_var: Decimal = Decimal(20)) -> tuple[list[dict], list[dict]]:
    uy, l, kdvsiz = [], [], []
    for r in kayitlar(yol, SUTUN, ("kod", "fiyat", "maliyet")):
        kod = metin(r["kod"])
        if not kod:
            continue
        d = {a: para(r[a]) if metin(r[a]) else None for a in SAYI}
        if d["fiyat"] is None or d["fiyat"] <= 0 or d["maliyet"] is None or any(metin(r[a]) and (d[a] is None or d[a] < 0) for a in SAYI) or (d["indirim"] or SIFIR) >= 100 or (d["komisyon"] or SIFIR) >= 100:
            uy.append(uyari("Yüksek", "Satır okunamadı", kod, f"Satır {r['_satir']}: fiyat, maliyet ya da oranlardan biri okunamadı (indirim ve komisyon 100'den küçük olmalı). Ürün alınmadı."))
            continue
        if d["kdv"] is None:
            kdvsiz.append(kod)
            d["kdv"] = kdv_var
        for a in ("komisyon", "kargo", "hizmet", "indirim", "kupon"):
            d[a] = d[a] or SIFIR
        l.append({"kod": kod, "ad": metin(r["ad"]), **d})
    if kdvsiz:
        uy.append(uyari("Bilgi", "KDV oranı yazılmamış", ", ".join(kdvsiz[:10]) + (" …" if len(kdvsiz) > 10 else ""), f"{len(kdvsiz)} üründe KDV boş; %{kdv_var} kullanıldı."))
    return l, uy


def birim(u: dict, fiyat: Decimal, matrah: str = "dahil") -> dict:
    net = fiyat / (1 + u["kdv"] / 100)
    komisyon = (fiyat if matrah == "dahil" else net) * u["komisyon"] / 100
    kar = net - komisyon - u["kargo"] - u["hizmet"] - u["maliyet"]
    return {"fiyat": fiyat, "net": net, "komisyon": komisyon, "kar": kar, "marj": kar * 100 / net if net else None}


def esik_fiyat(u: dict, marj: Decimal, matrah: str = "dahil") -> Decimal | None:
    """Marjın verilen yüzde olduğu müşteri fiyatı (KDV dahil)."""
    k = 1 + u["kdv"] / 100
    katsayi = (1 - marj / 100) / k - (u["komisyon"] / 100 if matrah == "dahil" else u["komisyon"] / 100 / k)
    sabit = u["kargo"] + u["hizmet"] + u["maliyet"]
    return sabit / katsayi if katsayi > 0 else None


def simule(urunler: list[dict], min_marj: Decimal = Decimal(10), matrah: str = "dahil") -> dict:
    if not urunler:
        raise ValueError("Simüle edilecek ürün yok.")
    uy, sonuc = [], []
    for u in urunler:
        kfiyat = u["fiyat"] * (1 - u["indirim"] / 100) - u["kupon"]
        if kfiyat <= 0:
            uy.append(uyari("Yüksek", "Kampanya fiyatı sıfır ya da eksi", u["kod"], f"Fiyat {u['fiyat']:,.2f}, indirim %{u['indirim']}, kupon {u['kupon']:,.2f}. Ürün hesaplanmadı."))
            continue
        m, k = birim(u, u["fiyat"], matrah), birim(u, kfiyat, matrah)
        s = {"urun": u, "mevcut": m, "kampanya": k, "indirim_tutari": u["fiyat"] - kfiyat, "indirim_orani": (u["fiyat"] - kfiyat) * 100 / u["fiyat"],
             "basabas": esik_fiyat(u, SIFIR, matrah), "min_marj_fiyat": esik_fiyat(u, min_marj, matrah),
             "gerekli_artis": (m["kar"] / k["kar"] - 1) * 100 if k["kar"] > 0 and m["kar"] > 0 else None,
             "mevcut_toplam": m["kar"] * u["adet"] if u["adet"] is not None else None, "kampanya_toplam": k["kar"] * u["beklenen"] if u["beklenen"] is not None else None}
        s["fark"] = s["kampanya_toplam"] - s["mevcut_toplam"] if s["mevcut_toplam"] is not None and s["kampanya_toplam"] is not None else None
        ad = u["ad"] or u["kod"]
        if k["kar"] < 0:
            uy.append(uyari("Yüksek", "Kampanyada zarar", u["kod"], f"{ad}: kampanya fiyatı {kfiyat:,.2f}, birim kâr {yuvarla(k['kar']):,.2f}. Başabaş fiyat {yuvarla(s['basabas']):,.2f}." if s["basabas"] else f"{ad}: kampanya birim kârı {yuvarla(k['kar']):,.2f}."))
        elif k["marj"] < min_marj:
            uy.append(uyari("Orta", "Marj eşik altında", u["kod"], f"{ad}: kampanya marjı %{yuvarla(k['marj'])}, eşik %{min_marj}."))
        if m["kar"] < 0:
            uy.append(uyari("Yüksek", "Mevcut fiyatta zarar", u["kod"], f"{ad}: kampanyasız birim kâr {yuvarla(m['kar']):,.2f}."))
        if s["fark"] is not None and s["fark"] < 0 and k["kar"] >= 0:
            uy.append(uyari("Orta", "Adet artışı yetmiyor", u["kod"], f"{ad}: beklenen {u['beklenen']:,.0f} adetle toplam kâr {yuvarla(s['kampanya_toplam']):,.2f}, mevcut {yuvarla(s['mevcut_toplam']):,.2f}"
                                                                      + (f"; kârı korumak için adet %{yuvarla(s['gerekli_artis'])} artmalı." if s["gerekli_artis"] is not None else ".")))
        sonuc.append(s)
    top = lambda a: sum((s[a] for s in sonuc if s[a] is not None), SIFIR) if any(s[a] is not None for s in sonuc) else None   # noqa: E731
    return {"urunler": sonuc, "toplam": {"mevcut": top("mevcut_toplam"), "kampanya": top("kampanya_toplam"), "fark": top("fark")}, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, min_marj: Decimal, matrah: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Simülasyon"
    y = lambda x: yuvarla(x) if x is not None else None   # noqa: E731
    _baslik(ws, ["Ürün Kodu", "Ürün Adı", "Fiyat", "Kampanya Fiyatı", "Toplam İndirim %", "Mevcut Birim Kâr", "Mevcut Marj %", "Kampanya Birim Kâr", "Kampanya Marjı %", "Gerekli Adet Artışı %",
                "Başabaş Fiyat", f"Min. Marj (%{min_marj}) Fiyatı", "Mevcut Adet", "Beklenen Adet", "Mevcut Toplam Kâr", "Kampanya Toplam Kâr", "Fark"], (14, 28) + (11,) * 15, "C2")
    for s in r["urunler"]:
        u, m, k = s["urun"], s["mevcut"], s["kampanya"]
        renk = KIRMIZI if k["kar"] < 0 else SARI if k["marj"] < min_marj else YESIL
        _satir(ws, [u["kod"], u["ad"], u["fiyat"], y(k["fiyat"]), y(s["indirim_orani"]), y(m["kar"]), y(m["marj"]), y(k["kar"]), y(k["marj"]), y(s["gerekli_artis"]), y(s["basabas"]), y(s["min_marj_fiyat"]),
                    u["adet"], u["beklenen"], y(s["mevcut_toplam"]), y(s["kampanya_toplam"]), y(s["fark"])],
               para_s=(3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 15, 16, 17), dolgu={9: renk, 17: None if s["fark"] is None else KIRMIZI if s["fark"] < 0 else YESIL})
    t = r["toplam"]
    if t["mevcut"] is not None or t["kampanya"] is not None:
        k = _satir(ws, ["TOPLAM", None] + [None] * 12 + [y(t["mevcut"]), y(t["kampanya"]), y(t["fark"])], para_s=(15, 16, 17))
        for h in ws[k]:
            h.font = Font(bold=True)
    ws.append([])
    ws.append([f"Gider kalemleri KDV hariç. Komisyon matrahı: müşteri fiyatı ({'KDV dahil' if matrah == 'dahil' else 'KDV hariç'}). Kupon satıcı karşılıyor varsayıldı. Gerekli adet artışı: toplam kârı korumak için."])
    ws = wb.create_sheet("Birim Kırılım")
    _baslik(ws, ["Ürün Kodu", "Senaryo", "Müşteri Fiyatı (KDV dahil)", "Net Satış (KDV hariç)", "Komisyon", "Kargo", "Hizmet Bedeli", "Birim Maliyet", "Birim Kâr", "Marj %"], (14, 10, 13, 13, 11, 10, 10, 11, 11, 9), "C2")
    for s in r["urunler"]:
        u = s["urun"]
        for ad in ("mevcut", "kampanya"):
            b = s[ad]
            _satir(ws, [u["kod"], ad.capitalize(), y(b["fiyat"]), y(b["net"]), y(b["komisyon"]), u["kargo"], u["hizmet"], u["maliyet"], y(b["kar"]), y(b["marj"])], para_s=range(3, 11),
                   dolgu={9: KIRMIZI if b["kar"] < 0 else None})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, min_marj: Decimal = Decimal(10), matrah: str = "dahil", kdv_var: Decimal = Decimal(20)) -> dict:
    urunler, uy = oku(yol, kdv_var)
    r = simule(urunler, min_marj, matrah)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, min_marj, matrah)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Planlanan indirim ve kupon kurgularının ürün bazında kâr marjına etkisini hesaplar.")
    p.add_argument("--urunler", type=Path, help="Ürün Kodu, Ürün Adı, Satış Fiyatı, KDV, Birim Maliyet, Komisyon, Kargo, Hizmet Bedeli, Mevcut Adet, İndirim, Kupon, Beklenen Adet")
    p.add_argument("--min-marj", default="10", help="Kampanya marjı bu yüzdenin altındaysa uyarılır (varsayılan 10)")
    p.add_argument("--matrah", choices=("dahil", "haric"), default="dahil", help="Komisyon KDV dahil müşteri fiyatına mı (varsayılan) KDV hariç tutara mı uygulanır")
    p.add_argument("--kdv", default="20", help="KDV sütunu boş olan ürünlerde kullanılacak oran (varsayılan 20)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kampanya_fiyat_simulasyonu.xlsx")
    a = p.parse_args(argv)
    yol = a.urunler or ORNEK / "kampanya_urunleri.csv"
    marj, kdv = para(a.min_marj), para(a.kdv)
    if marj is None or not -100 < marj < 100 or kdv is None or kdv < 0:
        print("[X] --min-marj -100 ile 100 arası, --kdv sıfır ya da üstü olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, marj, a.matrah, kdv)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    zarar = sum(1 for s in r["urunler"] if s["kampanya"]["kar"] < 0)
    dusuk = sum(1 for s in r["urunler"] if 0 <= s["kampanya"]["kar"] and s["kampanya"]["marj"] < marj)
    t = r["toplam"]
    print(f"[OK] {len(r['urunler'])} ürün · kampanyada zarar {zarar} · marjı %{marj} altında {dusuk}"
          + (f" · toplam kâr mevcut {t['mevcut']:,.2f} → kampanya {t['kampanya']:,.2f}" if t["mevcut"] is not None and t["kampanya"] is not None else ""))
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
