import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def u(kod, fiyat, kdv, maliyet, komisyon, kargo, hizmet=0, adet=None, indirim=0, kupon=0, beklenen=None):
    c = lambda x: None if x is None else D(x)   # noqa: E731
    return {"kod": kod, "ad": "", "fiyat": D(fiyat), "kdv": D(kdv), "maliyet": D(maliyet), "komisyon": D(komisyon), "kargo": D(kargo), "hizmet": D(hizmet), "adet": c(adet),
            "indirim": D(indirim), "kupon": D(kupon), "beklenen": c(beklenen)}


def veri():
    return [u("A", 1200, 20, 500, 15, 50, 10, 100, 10, 80, 180), u("B", 100, 10, 70, 20, 20), u("C", 100, 20, 10, 0, 0, kupon=100)]


def y(x):
    return main.yuvarla(x) if x is not None else None


class SimuleTesti(unittest.TestCase):
    def test_hesap(self):
        r = main.simule(veri())
        self.assertEqual([(s["urun"]["kod"], y(s["mevcut"]["kar"]), y(s["mevcut"]["marj"]), y(s["kampanya"]["fiyat"]), y(s["kampanya"]["net"]), y(s["kampanya"]["komisyon"]), y(s["kampanya"]["kar"]),
                           y(s["kampanya"]["marj"]), y(s["indirim_orani"]), y(s["gerekli_artis"]), y(s["basabas"]), y(s["min_marj_fiyat"]), y(s["mevcut_toplam"]), y(s["kampanya_toplam"]), y(s["fark"]))
                          for s in r["urunler"]],
                         [("A", D("260.00"), D("26.00"), D("1000.00"), D("833.33"), D("150.00"), D("123.33"), D("14.80"), D("16.67"), D("110.81"), D("819.51"), D("933.33"), D("26000.00"), D("22200.00"), D("-3800.00")),
                          ("B", D("-19.09"), D("-21.00"), D("100.00"), D("90.91"), D("20.00"), D("-19.09"), D("-21.00"), D("0.00"), None, D("126.92"), D("145.59"), None, None, None)])
        self.assertEqual(sorted((x["onem"], x["tur"], x["kim"]) for x in r["uyarilar"]),
                         [("Orta", "Adet artışı yetmiyor", "A"), ("Yüksek", "Kampanya fiyatı sıfır ya da eksi", "C"), ("Yüksek", "Kampanyada zarar", "B"), ("Yüksek", "Mevcut fiyatta zarar", "B")])
        self.assertEqual({k: y(v) for k, v in r["toplam"].items()}, {"mevcut": D("26000.00"), "kampanya": D("22200.00"), "fark": D("-3800.00")})

    def test_secenekler(self):
        r = main.simule(veri()[:1], D(18), "haric")
        s = r["urunler"][0]
        self.assertEqual((y(s["kampanya"]["komisyon"]), y(s["kampanya"]["kar"]), y(s["basabas"]), [x["tur"] for x in r["uyarilar"]]),
                         (D("125.00"), D("148.33"), D("790.59"), ["Marj eşik altında", "Adet artışı yetmiyor"]))
        self.assertIsNone(main.esik_fiyat(u("X", 1, 0, 1, 99, 0), D(10)))
        self.assertEqual(main.simule([u("Y", 100, 0, 10, 0, 0, adet=1)])["toplam"], {"mevcut": D(90), "kampanya": None, "fark": None})
        with self.assertRaises(ValueError):
            main.simule([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "kampanya_urunleri.csv", cikti)
            self.assertEqual([(s["urun"]["kod"], y(s["kampanya"]["kar"]), y(s["kampanya"]["marj"]), y(s["fark"])) for s in r["urunler"]],
                             [("EV-101", D("160.17"), D("9.50"), D("-26124.17")), ("EV-102", D("48.45"), D("8.08"), D("-31044.60")), ("EV-103", D("-71.43"), D("-9.27"), D("-46269.08")),
                              ("TK-201", D("58.41"), D("7.99"), D("-17361.88")), ("TK-202", D("64.02"), D("15.01"), D("2432.77")), ("KZ-301", D("-11.51"), D("-5.77"), D("-18026.00")),
                              ("KZ-302", D("183.39"), D("13.59"), D("3627.20")), ("KZ-303", D("34.79"), D("8.32"), None)])
            self.assertEqual([(x["tur"], x["kim"]) for x in r["uyarilar"]],
                             [("Kampanyada zarar", "EV-103"), ("Kampanyada zarar", "KZ-301"), ("Adet artışı yetmiyor", "EV-101"), ("Adet artışı yetmiyor", "EV-102"), ("Adet artışı yetmiyor", "TK-201"),
                              ("Marj eşik altında", "EV-101"), ("Marj eşik altında", "EV-102"), ("Marj eşik altında", "KZ-303"), ("Marj eşik altında", "TK-201"), ("KDV oranı yazılmamış", "KZ-303")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Simülasyon", "Birim Kırılım", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "u.csv"
            p.write_text("SKU;Fiyat;Maliyet;KDV Oranı;Komisyon Oranı\nA;1.250,00;500;;15\nB;0;1;;\nC;10;x;;\nD;10;1;;100\nE;10;1;-1;\n;1;1;;\n", encoding="utf-8")
            l, uy = main.oku(p, D(10))
            self.assertEqual(([(x["kod"], x["fiyat"], x["kdv"], x["komisyon"], x["kargo"], x["adet"]) for x in l], [x["tur"] for x in uy]),
                             ([("A", D(1250), D(10), D(15), D(0), None)], ["Satır okunamadı"] * 4 + ["KDV oranı yazılmamış"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--urunler", str(p), "--cikti", c]
            self.assertEqual(main.main(arg + ["--min-marj", "12,5", "--matrah", "haric", "--kdv", "10"]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--min-marj", "x"], ["--min-marj", "100"], ["--kdv", "-1"])], [2] * 3)
            self.assertEqual(main.main(["--urunler", str(Path(tmp) / "yok.csv")]), 1)
            p.write_text("SKU;Fiyat\nA;5\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
