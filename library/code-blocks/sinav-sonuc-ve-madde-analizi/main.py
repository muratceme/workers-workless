"""
Sınav Sonuç ve Madde Analizi — Workers / Workless kod bloğu
Eğitim › Ölçme Değerlendirme ve Rehberlik › Ölçme Değerlendirme Uzmanı

Optik okuyucu veya cevap dökümünden ve cevap anahtarından:
  - Öğrenci doğru, yanlış, boş, net (D − Y / k) ve 100 üzerinden puan; genel ve sınıf sırası.
  - Kitapçık türleri (A/B…) farklı soru sırasındaysa anahtar dosyasındaki eşleştirmeyle ortak soru numarasına çevrilir;
    iptal edilen sorular herkese doğru sayılır veya değerlendirme dışı bırakılır.
  - Klasik test kuramına göre madde analizi: güçlük (p), ayırt edicilik (üst-alt %27 grup farkı, d), düzeltilmiş
    madde-toplam korelasyonu (r), çeldirici analizi; KR-20 güvenirlik, ortalama, standart sapma, ölçmenin standart
    hatası.
  - Kazanım (konu) bazında sınıf başarı yüzdeleri.
İnternete bağlanmaz.

Kullanım:
    python main.py                                         # örnek sınavla dener
    python main.py --cevaplar cevaplar.xlsx --anahtar anahtar.xlsx
    python main.py --cevaplar cevaplar.xlsx --anahtar anahtar.xlsx --yanlis-katsayi 3 --iptal-dogru
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import Counter, OrderedDict, defaultdict
from pathlib import Path
from statistics import fmean, pstdev

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def katla(s) -> str:
    s = kucuk(s)
    for a, b in zip("ıöüşğç", "iousgc"):
        s = s.replace(a, b)
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def secenek(x) -> str:
    """Öğrenci cevabı: 'A'..'E'; boş/geçersiz → '' ; birden fazla işaret ('AB', '*') → '*'."""
    s = str(x or "").strip().upper()
    if not s or s in ("-", "_", "."):
        return ""
    if len(s) == 1 and s.isalpha():
        return s
    return "*"


# ----------------------------------------------------------------------------
# Okuma
# ----------------------------------------------------------------------------

def anahtar_oku(yol: Path) -> tuple[dict[str, dict[int, tuple[int, str]]], dict[int, dict]]:
    """Dönüş: (kitapçık → {kitapçıktaki soru no → (ortak soru no, doğru cevap)}, ortak soru no → {kazanım, iptal}).

    Biçim 1 (uzun): Kitapçık, Soru, Cevap [, A Soru No, Kazanım, İptal]
    Biçim 2 (geniş): Kitapçık, 1, 2, 3 … (tek satırda cevaplar; eşleştirme yoksa sorular kitapçıkta aynı sıradadır)
    """
    s = tablo_oku(yol)
    kb = [katla(x) for x in s[0]]
    bul = lambda *a: next((kb.index(katla(x)) for x in a if katla(x) in kb), None)  # noqa: E731
    i_kit, i_soru, i_cev = bul("kitapçık", "kitapcik", "grup", "form"), bul("soru", "soru no", "madde"), bul("cevap", "doğru cevap", "anahtar")
    i_ortak, i_kaz, i_ipt = bul("a soru no", "ortak soru no", "a karşılığı", "asıl soru no"), bul("kazanım", "konu", "kazanim"), bul("iptal")
    anahtar: dict[str, dict[int, tuple[int, str]]] = defaultdict(dict)
    sorular: dict[int, dict] = {}
    if i_soru is not None and i_cev is not None:
        for r in s[1:]:
            al = lambda i: r[i] if i is not None and i < len(r) else None  # noqa: E731
            if al(i_soru) in (None, ""):
                continue
            kit = str(al(i_kit) or "A").strip().upper()
            no = int(float(al(i_soru)))
            ortak = int(float(al(i_ortak))) if al(i_ortak) not in (None, "") else no
            anahtar[kit][no] = (ortak, secenek(al(i_cev)))
            bilgi = sorular.setdefault(ortak, {"kazanim": "", "iptal": False})
            if al(i_kaz):
                bilgi["kazanim"] = str(al(i_kaz)).strip()
            if katla(al(i_ipt)) in ("evet", "e", "x", "1", "iptal"):
                bilgi["iptal"] = True
    else:
        sutun = [(i, int(float(b))) for i, b in enumerate(s[0]) if re.fullmatch(r"\d+(?:\.0)?", str(b).strip())]
        for r in s[1:]:
            kit = str(r[i_kit] if i_kit is not None else "A").strip().upper()
            for i, no in sutun:
                anahtar[kit][no] = (no, secenek(r[i] if i < len(r) else ""))
                sorular.setdefault(no, {"kazanim": "", "iptal": False})
    if not anahtar:
        raise SystemExit(f"{yol.name}: cevap anahtarı okunamadı.")
    return dict(anahtar), dict(sorted(sorular.items()))


def cevaplari_oku(yol: Path) -> list[dict]:
    s = tablo_oku(yol)
    kb = [katla(x) for x in s[0]]
    bul = lambda *a: next((kb.index(katla(x)) for x in a if katla(x) in kb), None)  # noqa: E731
    i_no, i_ad = bul("öğrenci no", "ogrenci no", "numara", "no"), bul("ad soyad", "adı soyadı", "öğrenci", "ad")
    i_sinif, i_kit = bul("sınıf", "şube", "sinif"), bul("kitapçık", "kitapcik", "grup", "form")
    i_cevap = bul("cevaplar", "cevap dizisi")
    sutun = [(i, int(float(b))) for i, b in enumerate(s[0]) if re.fullmatch(r"(?:s)?\d+(?:\.0)?", str(b).strip().lower())
             and i not in (i_no, i_ad, i_sinif, i_kit)]
    sutun = [(i, int(re.sub(r"\D", "", str(s[0][i])) or 0)) for i, _ in sutun]
    if not sutun and i_cevap is None:
        raise SystemExit(f"{yol.name}: soru sütunları (1, 2, 3 … veya S1, S2 …) ya da tek 'Cevaplar' dizisi sütunu gerekli.")
    ogrenciler = []
    for n, r in enumerate(s[1:], 1):
        al = lambda i: r[i] if i is not None and i < len(r) else None  # noqa: E731
        if i_cevap is not None:
            dizi = str(al(i_cevap) or "")
            cev = {j + 1: secenek(c if c != " " else "") for j, c in enumerate(dizi)}
        else:
            cev = {no: secenek(al(i)) for i, no in sutun}
        ogrenciler.append({"no": str(al(i_no) or n).strip(), "ad": str(al(i_ad) or "").strip(), "sinif": str(al(i_sinif) or "-").strip(),
                           "kitapcik": str(al(i_kit) or "A").strip().upper() or "A", "ham": cev})
    return ogrenciler


# ----------------------------------------------------------------------------
# Puanlama ve analiz
# ----------------------------------------------------------------------------

def puanla(ogrenciler, anahtar, sorular, yanlis_katsayi: float, iptal_dogru: bool) -> list[str]:
    uyarilar = []
    gecerli = [no for no, b in sorular.items() if not (b["iptal"] and not iptal_dogru)]
    for o in ogrenciler:
        a = anahtar.get(o["kitapcik"])
        if a is None:
            uyarilar.append(f"{o['no']}: '{o['kitapcik']}' kitapçığının anahtarı yok; A kullanıldı")
            a = anahtar.get("A") or next(iter(anahtar.values()))
        o["cevap"], o["dogru_mu"] = {}, {}
        for kno, (ortak, dogru) in a.items():
            c = o["ham"].get(kno, "")
            o["cevap"][ortak] = c
            if sorular[ortak]["iptal"]:
                o["dogru_mu"][ortak] = 1 if iptal_dogru else None
            else:
                o["dogru_mu"][ortak] = 1 if c and c == dogru else 0
        d = sum(1 for no in gecerli if o["dogru_mu"].get(no) == 1)
        b = sum(1 for no in gecerli if not sorular[no]["iptal"] and o["cevap"].get(no, "") == "")
        y = len(gecerli) - d - b
        o.update(d=d, y=y, b=b, net=d - (y / yanlis_katsayi if yanlis_katsayi else 0),
                 puan=max(0.0, (d - (y / yanlis_katsayi if yanlis_katsayi else 0)) / len(gecerli) * 100) if gecerli else 0.0)
    return uyarilar


def madde_analizi(ogrenciler, anahtar, sorular, grup_orani: float = 0.27) -> tuple[OrderedDict, dict]:
    maddeler = [no for no, b in sorular.items() if not b["iptal"]]
    n = len(ogrenciler)
    toplam = {o["no"]: sum(o["dogru_mu"].get(m) or 0 for m in maddeler) for o in ogrenciler}
    sirali = sorted(ogrenciler, key=lambda o: -toplam[o["no"]])
    g = max(1, round(n * grup_orani))
    ust, alt = sirali[:g], sirali[-g:]
    dogru_cevap = {}
    for kit, a in anahtar.items():
        for _, (ortak, c) in a.items():
            dogru_cevap.setdefault(ortak, c)
    sonuc: OrderedDict[int, dict] = OrderedDict()
    for m in maddeler:
        x = [o["dogru_mu"].get(m) or 0 for o in ogrenciler]
        p = fmean(x) if n else 0.0
        pu, pa = fmean(o["dogru_mu"].get(m) or 0 for o in ust), fmean(o["dogru_mu"].get(m) or 0 for o in alt)
        kalan = [toplam[o["no"]] - (o["dogru_mu"].get(m) or 0) for o in ogrenciler]       # düzeltilmiş toplam (madde hariç)
        r = korelasyon(x, kalan)
        secenekler = Counter(o["cevap"].get(m, "") for o in ogrenciler)
        celdirici = {}
        for s in sorted(k for k in secenekler if k not in ("", "*")) or []:
            celdirici[s] = (sum(1 for o in ust if o["cevap"].get(m) == s) / g, sum(1 for o in alt if o["cevap"].get(m) == s) / g)
        sonuc[m] = {"p": p, "d": pu - pa, "r": r, "ust": pu, "alt": pa, "secenekler": secenekler, "celdirici": celdirici,
                    "dogru": dogru_cevap.get(m, ""), "kazanim": sorular[m]["kazanim"], "yorum": yorumla(p, pu - pa)}
    k = len(maddeler)
    xs = list(toplam.values())
    varyans = pstdev(xs) ** 2 if n > 1 else 0.0
    kr20 = (k / (k - 1)) * (1 - sum(v["p"] * (1 - v["p"]) for v in sonuc.values()) / varyans) if k > 1 and varyans > 0 else None
    sd = math.sqrt(varyans)
    istat = {"n": n, "k": k, "ort": fmean(xs) if xs else 0.0, "sd": sd, "kr20": kr20,
             "sem": sd * math.sqrt(1 - kr20) if kr20 is not None and kr20 < 1 else None,
             "ort_p": fmean(v["p"] for v in sonuc.values()) if sonuc else None, "grup": g}
    return sonuc, istat


def korelasyon(x: list[float], y: list[float]) -> float | None:
    if len(x) < 2:
        return None
    mx, my = fmean(x), fmean(y)
    sxy = sum((a - mx) * (b - my) for a, b in zip(x, y))
    sx = math.sqrt(sum((a - mx) ** 2 for a in x))
    sy = math.sqrt(sum((b - my) ** 2 for b in y))
    return sxy / (sx * sy) if sx and sy else None


def yorumla(p: float, d: float) -> str:
    """Ebel ve Frisbie ölçütleri (ayırt edicilik) ve yaygın güçlük aralıkları."""
    guc = "çok zor" if p < 0.20 else "zor" if p < 0.40 else "orta" if p <= 0.60 else "kolay" if p <= 0.80 else "çok kolay"
    if d >= 0.40:
        ay = "çok iyi ayırt ediyor"
    elif d >= 0.30:
        ay = "iyi ayırt ediyor"
    elif d >= 0.20:
        ay = "düzeltilmeli"
    else:
        ay = "ayırt etmiyor — gözden geçirin veya çıkarın"
    return f"{guc}; {ay}"


def celdirici_notlari(m: dict, secenek_sayisi: int = 4) -> list[str]:
    notlar = []
    for s, (u, a) in m["celdirici"].items():
        if s == m["dogru"]:
            continue
        if u - a >= 0.10:
            notlar.append(f"{s} şıkkını üst grup daha çok seçmiş (anahtar hatası veya belirsizlik?)")
    secilmeyen = [s for s in "ABCDEFGH"[:secenek_sayisi] if s != m["dogru"] and m["secenekler"].get(s, 0) == 0]
    if secilmeyen:
        notlar.append(f"hiç seçilmeyen çeldirici: {', '.join(secilmeyen)}")
    return notlar


def calistir(cevap_yolu: Path, anahtar_yolu: Path, cikti: Path, yanlis_katsayi: float = 4.0, iptal_dogru: bool = False) -> dict:
    anahtar, sorular = anahtar_oku(anahtar_yolu)
    ogrenciler = cevaplari_oku(cevap_yolu)
    if not ogrenciler:
        raise SystemExit("Öğrenci cevabı yok.")
    uyarilar = puanla(ogrenciler, anahtar, sorular, yanlis_katsayi, iptal_dogru)
    maddeler, istat = madde_analizi(ogrenciler, anahtar, sorular)
    harfler = {c for o in ogrenciler for c in o["cevap"].values() if c not in ("", "*")} |         {c for a in anahtar.values() for _, c in a.values() if c}
    istat["secenek"] = max(4, max((ord(c) - 64 for c in harfler if "A" <= c <= "H"), default=4))
    for i, o in enumerate(sorted(ogrenciler, key=lambda o: (-o["net"], o["no"])), 1):
        o["sira"] = i
    sinif = defaultdict(list)
    for o in ogrenciler:
        sinif[o["sinif"]].append(o)
    for ss in sinif.values():
        for i, o in enumerate(sorted(ss, key=lambda o: (-o["net"], o["no"])), 1):
            o["sinif_sira"] = i
    _rapor(ogrenciler, maddeler, istat, sorular, sinif, uyarilar, yanlis_katsayi, iptal_dogru, cikti)
    return {"ogrenciler": ogrenciler, "maddeler": maddeler, "istat": istat, "uyarilar": uyarilar}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
YESIL = PatternFill("solid", fgColor="E3F5E1")


def _baslik(ws, satir=1):
    for c in ws[satir]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI


def _rapor(ogrenciler, maddeler, istat, sorular, sinif, uyarilar, yk, iptal_dogru, cikti):
    wb = Workbook()
    o = wb.active
    o.title = "Özet"
    o.append(["Sınav sonuç ve madde analizi"])
    o["A1"].font = Font(bold=True, size=12)
    tr = lambda x, b=2: "-" if x is None else f"{x:.{b}f}".replace(".", ",")  # noqa: E731
    for s in [("Öğrenci", istat["n"]), ("Değerlendirilen madde", istat["k"]),
              ("İptal edilen soru", ", ".join(str(n) for n, b in sorular.items() if b["iptal"]) or "-"),
              ("Ortalama doğru (ham)", tr(istat["ort"])), ("Standart sapma", tr(istat["sd"])),
              ("Ortalama güçlük (p)", tr(istat["ort_p"])), ("KR-20 güvenirlik", tr(istat["kr20"], 3)),
              ("Ölçmenin standart hatası", tr(istat["sem"])), ("Üst/alt grup büyüklüğü (%27)", istat["grup"]),
              ("Net hesabı", f"doğru − yanlış / {yk:g}" if yk else "yalnız doğru"),
              ("İptal edilen sorular", "herkese doğru sayıldı" if iptal_dogru else "değerlendirme dışı bırakıldı")]:
        o.append(list(s))
    o.append([])
    o.append(["Sınıf", "Öğrenci", "Ortalama Net", "Ortalama Puan", "En Yüksek Net", "En Düşük Net", "Std. Sapma (net)"])
    _baslik(o, o.max_row)
    for ad, ss in sorted(sinif.items()):
        nets = [x["net"] for x in ss]
        o.append([ad, len(ss), round(fmean(nets), 2), round(fmean(x["puan"] for x in ss), 2), round(max(nets), 2),
                  round(min(nets), 2), round(pstdev(nets), 2) if len(nets) > 1 else 0])
    if istat["kr20"] is not None and istat["kr20"] < 0.70:
        o.append(["Uyarı", "KR-20 0,70'in altında: test güvenirliği düşük; ayırt etmeyen maddeleri gözden geçirin"])
    for u in uyarilar:
        o.append(["Uyarı", u])
    o.column_dimensions["A"].width = 30
    for c in "BCDEFG":
        o.column_dimensions[c].width = 14

    s = wb.create_sheet("Öğrenci Sonuçları")
    s.append(["Sıra", "Sınıf Sırası", "Öğrenci No", "Ad Soyad", "Sınıf", "Kitapçık", "Doğru", "Yanlış", "Boş", "Net", "Puan (100)"])
    _baslik(s)
    for x in sorted(ogrenciler, key=lambda x: x["sira"]):
        s.append([x["sira"], x["sinif_sira"], x["no"], x["ad"], x["sinif"], x["kitapcik"], x["d"], x["y"], x["b"], round(x["net"], 2),
                  round(x["puan"], 2)])
    for j, w in enumerate((6, 10, 11, 22, 8, 9, 7, 7, 6, 8, 10), 1):
        s.column_dimensions[get_column_letter(j)].width = w
    s.freeze_panes = "E2"
    s.auto_filter.ref = s.dimensions

    m = wb.create_sheet("Madde Analizi")
    m.append(["Soru", "Kazanım", "Doğru Cevap", "Güçlük (p)", "Üst Grup p", "Alt Grup p", "Ayırt Edicilik (d)", "Madde-Toplam r",
              "Yorum", "Çeldirici Notu", "A", "B", "C", "D", "E", "Boş", "Çoklu"])
    _baslik(m)
    for no, v in maddeler.items():
        sec = v["secenekler"]
        m.append([no, v["kazanim"], v["dogru"], round(v["p"], 3), round(v["ust"], 3), round(v["alt"], 3), round(v["d"], 3),
                  round(v["r"], 3) if v["r"] is not None else None, v["yorum"], "; ".join(celdirici_notlari(v, istat["secenek"])),
                  sec.get("A", 0), sec.get("B", 0), sec.get("C", 0), sec.get("D", 0), sec.get("E", 0), sec.get("", 0), sec.get("*", 0)])
        c = m.cell(m.max_row, 7)
        c.fill = YESIL if v["d"] >= 0.30 else SARI if v["d"] >= 0.20 else KIRMIZI
        if celdirici_notlari(v, istat["secenek"]):
            m.cell(m.max_row, 10).fill = SARI
    for j, w in enumerate((6, 26, 8, 10, 10, 10, 12, 12, 38, 44, 5, 5, 5, 5, 5, 5, 6), 1):
        m.column_dimensions[get_column_letter(j)].width = w
    m.freeze_panes = "C2"
    m.auto_filter.ref = m.dimensions

    kz = wb.create_sheet("Kazanım Analizi")
    kazanimlar = OrderedDict()
    for no, v in maddeler.items():
        kazanimlar.setdefault(v["kazanim"] or "(belirtilmemiş)", []).append(no)
    siniflar = sorted(sinif)
    kz.append(["Kazanım", "Sorular", "Genel Başarı %"] + [f"{x} %" for x in siniflar])
    _baslik(kz)
    for kaz, nolar in kazanimlar.items():
        genel = fmean(fmean(o["dogru_mu"].get(n) or 0 for o in ogrenciler) for n in nolar)
        satir = [kaz, ", ".join(map(str, nolar)), round(genel * 100, 1)]
        for sn in siniflar:
            satir.append(round(fmean(fmean(o["dogru_mu"].get(n) or 0 for o in sinif[sn]) for n in nolar) * 100, 1))
        kz.append(satir)
        for j in range(3, len(satir) + 1):
            v = kz.cell(kz.max_row, j).value
            kz.cell(kz.max_row, j).fill = KIRMIZI if v < 50 else SARI if v < 70 else YESIL
    kz.column_dimensions["A"].width = 36
    kz.column_dimensions["B"].width = 16
    for j in range(3, len(siniflar) + 4):
        kz.column_dimensions[get_column_letter(j)].width = 12

    b = wb.create_sheet("Bilgi")
    for x in [["Güçlük (p)", "doğru cevaplayanların oranı: < 0,20 çok zor · 0,20–0,39 zor · 0,40–0,60 orta · 0,61–0,80 kolay · > 0,80 çok kolay"],
              ["Ayırt edicilik (d)", "üst %27 grubun p'si − alt %27 grubun p'si. Ebel: ≥ 0,40 çok iyi · 0,30–0,39 iyi · 0,20–0,29 düzeltilmeli · "
                                     "< 0,20 çıkarılmalı veya yeniden yazılmalı"],
              ["Madde-toplam r", "madde puanı ile maddenin çıkarıldığı toplam puan arasındaki korelasyon (düzeltilmiş nokta-çift serili)"],
              ["KR-20", "k/(k−1) × (1 − Σ p(1−p) / σ²); σ² toplam ham puanların varyansı. ≥ 0,70 genelde yeterli kabul edilir"],
              ["Çeldirici", "yanlış bir şıkkı üst grup alt gruptan daha çok seçiyorsa anahtar veya soru kökü kontrol edilmeli; hiç "
                            "seçilmeyen çeldirici işlevsizdir"],
              ["Net", "4 yanlışın 1 doğruyu götürmesi 5 seçenekli sınavlarda yaygındır; 4 seçenekli sınavlarda --yanlis-katsayi 3, "
                      "yanlış götürmeyen sınavlarda 0 kullanın"]]:
        b.append(x)
    b.column_dimensions["A"].width = 18
    b.column_dimensions["B"].width = 130
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ornek = BURASI / "ornek_veri"
    ap = argparse.ArgumentParser(description="Sınav cevaplarından net/puan ve klasik test kuramına göre madde analizi çıkarır.")
    ap.add_argument("--cevaplar", type=Path, default=ornek / "cevaplar.csv",
                    help="Cevaplar (.xlsx/.csv): Öğrenci No, Ad Soyad, Sınıf, Kitapçık, 1, 2, 3 … (veya tek 'Cevaplar' dizisi)")
    ap.add_argument("--anahtar", type=Path, default=ornek / "anahtar.csv",
                    help="Anahtar: Kitapçık, Soru, Cevap [, A Soru No, Kazanım, İptal] — veya Kitapçık, 1, 2, 3 …")
    ap.add_argument("--yanlis-katsayi", type=float, default=4.0, help="Kaç yanlış bir doğruyu götürür (varsayılan 4; 0 = götürmez)")
    ap.add_argument("--iptal-dogru", action="store_true", help="İptal edilen soruları herkese doğru say (varsayılan: değerlendirme dışı)")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "sinav_analizi.xlsx")
    a = ap.parse_args(argv)
    s = calistir(a.cevaplar, a.anahtar, a.cikti, a.yanlis_katsayi, a.iptal_dogru)
    i = s["istat"]
    zayif = [no for no, v in s["maddeler"].items() if v["d"] < 0.20]
    print(f"[OK] {i['n']} öğrenci · {i['k']} madde · ortalama {i['ort']:.2f} doğru · KR-20 "
          f"{'-' if i['kr20'] is None else format(i['kr20'], '.3f')} · ayırt etmeyen madde: {', '.join(map(str, zayif)) or 'yok'}")
    for u in s["uyarilar"]:
        print(f"[!] {u}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
