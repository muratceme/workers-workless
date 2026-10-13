import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
ENDEKS = {(2026, 1): D(100), (2026, 2): D(110), (2026, 3): D(121), (2026, 4): D(125)}
SAY = [0]


def a(ay, fiyat, miktar=10, kod="M1", pb="TL", tedarikci="T1", gun=10):
    SAY[0] += 1
    return {"tarih": date(2026, ay, gun), "kod": kod, "ad": f"Malzeme {kod}", "tedarikci": tedarikci, "miktar": None if miktar is None else D(miktar), "fiyat": D(str(fiyat)), "pb": pb, "satir": SAY[0]}


class AnalizTesti(unittest.TestCase):
    def test_degisim_endeks_ve_esik(self):
        r = main.analiz([a(1, 100), a(2, 120), a(3, 145.2), a(4, 150, tedarikci="T2")], ENDEKS)       # endeks: +%10, +%10, +%3,31
        self.assertEqual([(x["degisim"], x["endeks_degisim"], x["fark"], x["isaret"]) for x in r["satirlar"]],
                         [(None, None, None, ""), (D("20.00"), D("10.00"), D("10.00"), ""), (D("21.00"), D("10.00"), D("11.00"), "Olağandışı artış"), (D("3.31"), D("3.31"), D("0.00"), "")])
        m = r["malzemeler"][0]
        self.assertEqual((m["alim"], m["miktar"], m["tutar"], m["ortalama"], m["ilk"], m["son"], m["en_dusuk"], m["en_yuksek"], m["nominal"], m["endeks"], m["reel"], m["tedarikci"], m["tedarikci_sayisi"], m["olagandisi"]),
                         (4, D(40), D("5152.00"), D("128.8000"), D("100.0000"), D("150.0000"), D("100.0000"), D("150.0000"), D("50.00"), D("25.00"), D("20.00"), "T2", 2, 1))
        self.assertEqual([(u["tur"], u["kim"]) for u in r["uyarilar"]], [("Olağandışı fiyat artışı", "M1 · 10.03.2026")])
        self.assertEqual(main.analiz([a(1, 100), a(2, 120)], ENDEKS, esik=D("9.99"))["olagandisi"], 1)
        r = main.analiz([a(1, 100), a(2, 110), a(3, 121.1)])                                            # endeks yok: yalnız nominal değişim
        self.assertEqual(([(x["degisim"], x["fark"], x["isaret"]) for x in r["satirlar"]], r["malzemeler"][0]["reel"], r["endeks_var"]),
                         ([(None, None, ""), (D("10.00"), None, ""), (D("10.09"), None, "Olağandışı artış")], None, False))
        r = main.analiz([a(1, 100), a(5, 150), a(6, 150)], ENDEKS)
        self.assertEqual(([x["isaret"] for x in r["satirlar"]], sorted((u["tur"], u["kim"]) for u in r["uyarilar"]), r["malzemeler"][0]["reel"]),
                         (["", "Olağandışı artış", ""], [("Endeks eksik", "2 dönem"), ("Olağandışı fiyat artışı", "M1 · 10.05.2026")], None))

    def test_doviz_ve_gruplama(self):
        kur = {((2026, 1), "USD"): D(40), ((2026, 2), "USD"): D(50), ((2026, 3), "EUR"): D(60)}
        l = [a(2, 2.2, kod="D1", pb="USD"), a(1, 2, kod="D1", pb="USD"), a(3, 2, kod="D1", pb="EUR"), a(4, 3, kod="D1", pb="USD"), a(1, 5, None, kod="tek"), a(2, 9, kod="B", gun=5), a(2, 8, kod="B", gun=5), a(1, 10, kod="b")]
        r = main.analiz(l, None, kur, D(30))
        self.assertEqual([(x["kod"], x["tarih"].month, x["tl"], x["degisim"], x["doviz_degisim"], x["isaret"]) for x in r["satirlar"]],
                         [("b", 1, D("10.0000"), None, None, ""), ("B", 2, D("9.0000"), D("-10.00"), None, ""), ("B", 2, D("8.0000"), D("-11.11"), None, ""),
                          ("D1", 1, D("80.0000"), None, None, ""), ("D1", 2, D("110.0000"), D("37.50"), D("10.00"), "Olağandışı artış"), ("D1", 3, D("120.0000"), D("9.09"), None, ""), ("tek", 1, D("5.0000"), None, None, "")])
        self.assertEqual([(m["kod"], m["alim"], m["nominal"], m["miktar"], m["ortalama"]) for m in r["malzemeler"]],
                         [("D1", 3, D("50.00"), D(30), D("103.3333")), ("B", 3, D("-20.00"), D(30), D("9.0000")), ("tek", 1, None, None, None)])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]), [("Bilgi", "Tek alım", "tek"), ("Orta", "Olağandışı fiyat artışı", "D1 · 10.02.2026"), ("Yüksek", "Kur bulunamadı", "D1 · 10.04.2026")])
        self.assertIn("Döviz fiyatı %+10.00", next(u["aciklama"] for u in r["uyarilar"] if u["onem"] == "Orta"))
        with self.assertRaises(ValueError):
            main.analiz([a(1, 2, pb="USD")])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "f.xlsx"
            r = main.calistir(ORNEK / "alimlar.csv", cikti, ORNEK / "endeks.csv", ORNEK / "kurlar.csv")
            self.assertEqual([(m["kod"], m["alim"], m["tutar"], m["ortalama"], m["nominal"], m["endeks"], m["reel"], m["olagandisi"]) for m in r["malzemeler"]],
                             [("M-100", 4, D("215900.00"), D("4.5936"), D("35.00"), D("16.00"), D("16.38"), 1), ("M-300", 3, D("171000.00"), D("65.7692"), D("20.97"), D("14.56"), D("5.59"), 1),
                              ("M-200", 3, D("768900.00"), D("48.0563"), D("21.53"), D("21.00"), D("0.44"), 0), ("M-400", 1, D("17500.00"), D("0.3500"), None, None, None, 0)])
            self.assertEqual([(x["kod"], x["tl"], x["degisim"], x["doviz_degisim"], x["endeks_degisim"], x["fark"], x["isaret"]) for x in r["satirlar"] if x["degisim"] is not None],
                             [("M-100", D("4.2000"), D("5.00"), None, D("6.00"), D("-1.00"), ""), ("M-100", D("4.4500"), D("5.95"), None, D("4.72"), D("1.24"), ""),
                              ("M-100", D("5.4000"), D("21.35"), None, D("4.50"), D("16.84"), "Olağandışı artış"), ("M-200", D("47.5800"), D("10.14"), D("1.67"), D("11.00"), D("-0.86"), ""),
                              ("M-200", D("52.5000"), D("10.34"), D("2.46"), D("9.01"), D("1.33"), ""), ("M-300", D("58.0000"), D("-6.45"), None, D("9.71"), D("-16.16"), ""),
                              ("M-300", D("75.0000"), D("29.31"), None, D("4.42"), D("24.89"), "Olağandışı artış")])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Orta", "Olağandışı fiyat artışı", "M-100 · 18.07.2026"), ("Orta", "Olağandışı fiyat artışı", "M-300 · 22.08.2026"), ("Bilgi", "Tek alım", "M-400")])
            self.assertIn("Tedarikçi değişmiş: Ambalaj C Ltd. → Ambalaj A Ltd.", r["uyarilar"][1]["aciklama"])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Malzeme Özeti", "Alım Geçmişi", "Uyarılar"])

    def test_okuma_ve_cli(self):
        self.assertEqual([main.donem(x) for x in ("09.2026", "9/2026", "2026-09", "15.09.2026", date(2026, 9, 1), "13.2026", "x")], [(2026, 9)] * 5 + [None, None])
        with tempfile.TemporaryDirectory() as tmp:
            y, e, k = Path(tmp) / "a.csv", Path(tmp) / "e.csv", Path(tmp) / "k.csv"
            y.write_text("Fatura Tarihi;Stok Kodu;Stok Adı;Satıcı;Adet;Fiyat;Döviz\n10.01.2026;K1;Koli;T;1000;4.25;\n10.02.2026;K1;Koli;T;;1.250,5;try\n10.03.2026;K1;Koli;T;0;5;\nx;K1;Koli;T;1;5;\n10.03.2026;K1;Koli;T;1;-5;\n"
                         "10.03.2026;;Bant;T;2;3;eur\n", encoding="utf-8")
            e.write_text("Ay;ÜFE\n01.2026;100,5\n2026-02;x\n", encoding="utf-8")
            k.write_text("Ay;Döviz;Ortalama Kur\n03.2026;eur;50,25\n04.2026;;1\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["kod"], x["ad"], x["miktar"], x["fiyat"], x["pb"]) for x in l], [u["tur"] for u in uy]),
                             ([("K1", "Koli", D(1000), D("4.25"), "TL"), ("K1", "Koli", None, D("1250.5"), "TL"), ("", "Bant", D(2), D(3), "EUR")], ["Alım okunamadı"] * 3))
            self.assertEqual((main.endeks_oku(e), main.kur_oku(k)), ({(2026, 1): D("100.5")}, {((2026, 3), "EUR"): D("50.25")}))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--alimlar", str(y), "--cikti", c]
            self.assertEqual(main.calistir(y, Path(c), e, k)["satirlar"][0]["tl"], D("150.7500"))
            self.assertEqual(main.main(arg + ["--endeks", str(e), "--kurlar", str(k), "--esik", "5"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual((main.main(arg + ["--esik", "x"]), main.main(arg + ["--esik=-1"])), (2, 2))
            self.assertEqual([main.main(x) for x in (["--alimlar", str(Path(tmp) / "yok.csv")], arg + ["--endeks", str(Path(tmp) / "yok.csv")], arg + ["--kurlar", str(Path(tmp) / "yok.csv")])], [1, 1, 1])
            for icerik in ("Tarih;Malzeme\n10.01.2026;A\n", "Tarih;Malzeme;Birim Fiyat;Para Birimi\n10.01.2026;A;5;USD\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
