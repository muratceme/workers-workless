"""
Satın Alma Fiyat Geçmişi Analizi — Gerisi Bende kod bloğu
Satın Alma › Satın Alma Uzmanı

Malzeme bazında geçmiş alım fiyatlarını kur ve fiyat endeksiyle birlikte inceler, olağandışı artışları işaretler.
  TL birim fiyat        = birim fiyat × alım ayının kuru (--kurlar; TL = 1)
  Önceki alıma göre %   = (TL fiyat ÷ aynı malzemenin önceki alımındaki TL fiyat − 1) × 100
  Endeks değişimi %     = (alım ayının endeksi ÷ önceki alım ayının endeksi − 1) × 100   (--endeks dosyasından)
  Fark (puan)           = fiyat değişimi % − endeks değişimi %
  Olağandışı artış      = fark > --esik (10 puan); endeks yoksa fiyat değişimi > --esik
  Nominal değişim %     = (son fiyat ÷ ilk fiyat − 1) × 100;   reel değişim % = ((1 + nominal) ÷ (1 + endeks değişimi) − 1) × 100
  Ağırlıklı ort. fiyat  = Σ(miktar × TL fiyat) ÷ Σ miktar
Endeksi ve kurları siz verirsiniz; kod veri çekmez. İşaret bir inceleme önerisidir, haksız fiyat tespiti değildir.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
SUTUN = {"tarih": ("tarih", "alim tarihi", "fatura tarihi", "siparis tarihi"), "kod": ("malzeme kodu", "stok kodu", "kod"), "ad": ("malzeme", "malzeme adi", "stok adi", "urun"), "tedarikci": ("tedarikci", "satici", "firma"),
         "miktar": ("miktar", "adet"), "fiyat": ("birim fiyat", "fiyat", "birim fiyati"), "pb": ("para birimi", "doviz", "doviz cinsi", "pb")}
ENDEKS = {"donem": ("donem", "ay", "tarih"), "endeks": ("endeks", "endeks degeri", "ufe", "tufe")}
KUR = {"donem": ("donem", "ay", "tarih"), "pb": ("para birimi", "doviz", "doviz cinsi", "pb"), "kur": ("kur", "ortalama kur", "doviz kuru")}
Z, YUZ, BIR = Decimal(0), Decimal(100), Decimal(1)
FF = "#,##0.00##"


def olcu(x) -> Decimal | None:
    """Miktar ve birim fiyat: virgül veya tek nokta ondalık ayırıcıdır (4.25 = 4,25; 1.250,5 = 1250,5)."""
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x or ""))
    s = s.replace(".", "").replace(",", ".") if "," in s else s
    try:
        return Decimal(s) if s else None
    except InvalidOperation:
        return None


def donem(x) -> tuple[int, int] | None:
    if isinstance(x, (date, datetime)):
        return x.year, x.month
    s = metin(x)
    m = re.fullmatch(r"(\d{1,2})[./-](\d{4})", s)
    n = re.fullmatch(r"(\d{4})[./-](\d{1,2})", s)
    y, a = (int(m.group(2)), int(m.group(1))) if m else (int(n.group(1)), int(n.group(2))) if n else (0, 0)
    if 1 <= a <= 12:
        return y, a
    t = tarih(s)
    return (t.year, t.month) if t else None


def pb_duzelt(x) -> str:
    s = metin(x).upper().replace("₺", "TL")
    return "TL" if s in ("", "TL", "TRY", "TRL") else s


def endeks_oku(yol: Path) -> dict:
    return {d: v for r in kayitlar(yol, ENDEKS, ("donem", "endeks")) for d, v in [(donem(r["donem"]), olcu(r["endeks"]))] if d and v and v > 0}


def kur_oku(yol: Path) -> dict:
    return {(d, pb_duzelt(r["pb"])): v for r in kayitlar(yol, KUR, ("donem", "pb", "kur")) for d, v in [(donem(r["donem"]), olcu(r["kur"]))] if d and v and v > 0 and metin(r["pb"])}


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("tarih", "ad", "fiyat")):
        ad, kod = metin(r["ad"]), metin(r["kod"])
        if not ad and not kod:
            continue
        t, f, m = tarih(r["tarih"]), olcu(r["fiyat"]), olcu(r["miktar"])
        if t is None or f is None or f <= 0 or (metin(r["miktar"]) and (m is None or m <= 0)):
            uy.append(uyari("Yüksek", "Alım okunamadı", f"{kod or ad} (satır {r['_satir']})", f"Tarih '{metin(r['tarih'])}', miktar '{metin(r['miktar'])}', birim fiyat '{metin(r['fiyat'])}'. Fiyat ve miktar pozitif olmalı. Satır analize alınmadı."))
            continue
        l.append({"tarih": t, "kod": kod, "ad": ad or kod, "tedarikci": metin(r["tedarikci"]), "miktar": m, "fiyat": f, "pb": pb_duzelt(r["pb"]), "satir": r["_satir"]})
    return l, uy


def yuzde(yeni: Decimal, eski: Decimal) -> Decimal:
    return (yeni / eski - 1) * YUZ


def analiz(alimlar: list[dict], endeks: dict | None = None, kurlar: dict | None = None, esik: Decimal = Decimal(10)) -> dict:
    endeks, kurlar, uy, grup = endeks or {}, kurlar or {}, [], defaultdict(list)
    for a in sorted(alimlar, key=lambda a: (a["tarih"], a["satir"])):
        ay = (a["tarih"].year, a["tarih"].month)
        kur = BIR if a["pb"] == "TL" else kurlar.get((ay, a["pb"]))
        if kur is None:
            uy.append(uyari("Yüksek", "Kur bulunamadı", f"{a['kod'] or a['ad']} · {a['tarih']:%d.%m.%Y}", f"{ay[1]:02d}.{ay[0]} dönemi için {a['pb']} kuru verilmedi; alım analize alınmadı. --kurlar dosyasına ekleyin."))
            continue
        a["ay"], a["kur"], a["tl"] = ay, kur, yuvarla(a["fiyat"] * kur, "0.0001")
        grup[katla(a["kod"]) or katla(a["ad"])].append(a)
    if not grup:
        raise ValueError("Analiz edilecek alım yok.")
    satirlar, malzemeler, eksik_endeks = [], [], set()
    for liste in grup.values():
        olagandisi = 0
        for i, a in enumerate(liste):
            a.update(degisim=None, endeks_degisim=None, fark=None, doviz_degisim=None, isaret="")
            if i == 0:
                continue
            o = liste[i - 1]
            deg = yuzde(a["tl"], o["tl"])
            a["degisim"] = yuvarla(deg)
            if a["pb"] != "TL" and a["pb"] == o["pb"]:
                a["doviz_degisim"] = yuvarla(yuzde(a["fiyat"], o["fiyat"]))
            if a["ay"] in endeks and o["ay"] in endeks:
                ed = yuzde(endeks[a["ay"]], endeks[o["ay"]])
                a["endeks_degisim"], a["fark"] = yuvarla(ed), yuvarla(deg - ed)
            elif endeks:
                eksik_endeks |= {d for d in (a["ay"], o["ay"]) if d not in endeks}
            olcut = deg - ed if a["fark"] is not None else deg
            if olcut > esik:
                a["isaret"], olagandisi = "Olağandışı artış", olagandisi + 1
                uy.append(uyari("Orta", "Olağandışı fiyat artışı", f"{a['kod'] or a['ad']} · {a['tarih']:%d.%m.%Y}",
                                f"{a['ad']}: {o['tl']:,.4f} → {a['tl']:,.4f} TL (%{a['degisim']:+}); " + (f"aynı dönemde endeks %{a['endeks_degisim']:+}, fark {a['fark']:+} puan." if a["fark"] is not None else "endeks verisi yok.")
                                + (f" Döviz fiyatı %{a['doviz_degisim']:+}; kalan kısım kur etkisi." if a["doviz_degisim"] is not None else "")
                                + (f" Tedarikçi değişmiş: {o['tedarikci']} → {a['tedarikci']}" if katla(o["tedarikci"]) != katla(a["tedarikci"]) else f" Tedarikçi: {a['tedarikci']}").rstrip(".") + "."))
        ilk, son = liste[0], liste[-1]
        miktarli = [a for a in liste if a["miktar"]]
        top_miktar, top_tutar = sum((a["miktar"] for a in miktarli), Z), sum((a["miktar"] * a["tl"] for a in miktarli), Z)
        m = {"kod": son["kod"], "ad": son["ad"], "alim": len(liste), "miktar": top_miktar or None, "tutar": yuvarla(top_tutar) if miktarli else None, "ortalama": yuvarla(top_tutar / top_miktar, "0.0001") if top_miktar else None,
             "ilk_tarih": ilk["tarih"], "ilk": ilk["tl"], "son_tarih": son["tarih"], "son": son["tl"], "en_dusuk": min(a["tl"] for a in liste), "en_yuksek": max(a["tl"] for a in liste), "nominal": None, "endeks": None, "reel": None,
             "tedarikci": son["tedarikci"], "tedarikci_sayisi": len({katla(a["tedarikci"]) for a in liste if a["tedarikci"]}), "olagandisi": olagandisi}
        if len(liste) > 1:
            nom = yuzde(son["tl"], ilk["tl"])
            m["nominal"] = yuvarla(nom)
            if ilk["ay"] in endeks and son["ay"] in endeks:
                ed = yuzde(endeks[son["ay"]], endeks[ilk["ay"]])
                m["endeks"], m["reel"] = yuvarla(ed), yuvarla(((BIR + nom / YUZ) / (BIR + ed / YUZ) - 1) * YUZ)
        else:
            uy.append(uyari("Bilgi", "Tek alım", m["kod"] or m["ad"], "Malzemenin tek alımı var; fiyat değişimi hesaplanamadı."))
        malzemeler.append(m)
        satirlar += liste
    if eksik_endeks:
        uy.append(uyari("Orta", "Endeks eksik", f"{len(eksik_endeks)} dönem", "Şu dönemler için endeks verilmedi: " + ", ".join(f"{a:02d}.{y}" for y, a in sorted(eksik_endeks)) + ". Bu alımlarda yalnız fiyat değişimi eşikle karşılaştırıldı."))
    malzemeler.sort(key=lambda m: (-m["olagandisi"], -(m["reel"] if m["reel"] is not None else m["nominal"] if m["nominal"] is not None else Decimal(-10 ** 9)), katla(m["kod"] or m["ad"])))
    return {"satirlar": sorted(satirlar, key=lambda a: (katla(a["kod"]) or katla(a["ad"]), a["tarih"], a["satir"])), "malzemeler": malzemeler, "olagandisi": sum(m["olagandisi"] for m in malzemeler), "endeks_var": bool(endeks), "esik": esik,
            "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Malzeme Özeti"
    _baslik(ws, ["Malzeme Kodu", "Malzeme", "Alım Sayısı", "Toplam Miktar", "Toplam Tutar (TL)", "Ağırlıklı Ort. Fiyat", "İlk Alım", "İlk Fiyat", "Son Alım", "Son Fiyat", "En Düşük", "En Yüksek", "Nominal Değişim %", "Endeks Değişimi %",
                 "Reel Değişim %", "Son Tedarikçi", "Tedarikçi Sayısı", "Olağandışı Artış"], (13, 30, 8, 13, 16, 12, 12, 11, 12, 11, 11, 11, 11, 11, 11, 24, 9, 10))
    for m in r["malzemeler"]:
        n = _satir(ws, [m["kod"], m["ad"], m["alim"], m["miktar"], m["tutar"], m["ortalama"], m["ilk_tarih"], m["ilk"], m["son_tarih"], m["son"], m["en_dusuk"], m["en_yuksek"], m["nominal"], m["endeks"], m["reel"], m["tedarikci"],
                        m["tedarikci_sayisi"], m["olagandisi"] or None], para_s=(4, 5, 13, 14, 15), tarih_s=(7, 9), dolgu={18: SARI if m["olagandisi"] else None})
        for j in (6, 8, 10, 11, 12):
            ws.cell(n, j).number_format = FF
    ws.append([])
    ws.append(["Fiyatlar TL karşılığıdır. Reel değişim, verdiğiniz endekse göre arındırılmış değişimdir." if r["endeks_var"] else "Endeks verilmedi; reel değişim hesaplanmadı, eşik yalnız nominal değişime uygulandı."])
    ws = wb.create_sheet("Alım Geçmişi")
    _baslik(ws, ["Malzeme Kodu", "Malzeme", "Tarih", "Tedarikçi", "Miktar", "Birim Fiyat", "Para Birimi", "Kur", "TL Birim Fiyat", "Önceki Alıma Göre %", "Döviz Fiyatı Değişim %", "Endeks Değişimi %", "Fark (puan)", "İşaret"],
            (13, 30, 12, 24, 12, 11, 8, 9, 12, 11, 11, 11, 10, 16))
    for a in r["satirlar"]:
        n = _satir(ws, [a["kod"], a["ad"], a["tarih"], a["tedarikci"], a["miktar"], a["fiyat"], a["pb"], a["kur"] if a["pb"] != "TL" else None, a["tl"], a["degisim"], a["doviz_degisim"], a["endeks_degisim"], a["fark"], a["isaret"]],
                   para_s=(5, 10, 11, 12, 13), tarih_s=(3,), dolgu={14: SARI if a["isaret"] else None})
        for j in (6, 8, 9):
            ws.cell(n, j).number_format = FF
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, endeks_yol: Path | None = None, kur_yol: Path | None = None, esik: Decimal = Decimal(10)) -> dict:
    alimlar, uy = oku(yol)
    r = analiz(alimlar, endeks_oku(endeks_yol) if endeks_yol else None, kur_oku(kur_yol) if kur_yol else None, esik)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Malzeme bazında alım fiyat geçmişini kur ve endeksle birlikte analiz eder.")
    p.add_argument("--alimlar", type=Path, help="Tarih, Malzeme Kodu, Malzeme, Tedarikçi, Miktar, Birim Fiyat, Para Birimi")
    p.add_argument("--endeks", type=Path, help="Dönem (AA.YYYY), Endeks — karşılaştırmada kullanılacak aylık fiyat endeksi")
    p.add_argument("--kurlar", type=Path, help="Dönem (AA.YYYY), Para Birimi, Kur — dövizli alımlar için aylık kur")
    p.add_argument("--esik", default="10", help="Olağandışı artış eşiği, puan (varsayılan 10)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "fiyat_gecmisi_analizi.xlsx")
    a = p.parse_args(argv)
    ornek = a.alimlar is None
    yol, endeks, kurlar = (ORNEK / "alimlar.csv", a.endeks or ORNEK / "endeks.csv", a.kurlar or ORNEK / "kurlar.csv") if ornek else (a.alimlar, a.endeks, a.kurlar)
    esik = olcu(a.esik)
    if esik is None or esik < 0:
        print("[X] --esik sıfır veya pozitif sayı olmalı")
        return 2
    for y in (yol, endeks, kurlar):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(yol, a.cikti, endeks, kurlar, esik)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(r['malzemeler'])} malzeme, {len(r['satirlar'])} alım · {r['olagandisi']} olağandışı artış" + ("" if r["endeks_var"] else " · endeks verilmedi, yalnız nominal değişime bakıldı"))
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek" or u["tur"] == "Olağandışı fiyat artışı":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
