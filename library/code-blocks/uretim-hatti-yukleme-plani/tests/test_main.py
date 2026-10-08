import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BAS = date(2026, 10, 12)


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "h.xlsx"
        cls.s = main.calistir(ORNEK / "hatlar.csv", ORNEK / "siparisler.csv", cls.cikti, BAS)
        cls.p = {x.no: x for x in cls.s["siparisler"]}
        cls.u = {}
        for u in cls.s["uyarilar"]:
            cls.u.setdefault(u["tur"], []).append(u)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_kapasite_ve_ogrenme(self):
        self.assertAlmostEqual(self.s["hatlar"]["H1"].gunluk_dk, 22 * 540 * 0.65)
        g = self.p["PO-501"].gunluk
        self.assertAlmostEqual(g[date(2026, 10, 12)], 7722 * 0.5 / 12)
        self.assertAlmostEqual(g[date(2026, 10, 13)], 7722 * 0.7 / 12)
        self.assertAlmostEqual(g[date(2026, 10, 15)], 7722 / 12)
        self.assertAlmostEqual(g[date(2026, 10, 28)], 7722 / 24)          # arife yarım gün
        self.assertNotIn(date(2026, 10, 29), g)                           # Cumhuriyet Bayramı
        self.assertNotIn(date(2026, 10, 18), g)                           # Pazar
        self.assertAlmostEqual(sum(g.values()), 12000)

    def test_atamalar(self):
        self.assertEqual({k: (x.hat, x.durum) for k, x in self.p.items()},
                         {"PO-501": ("H1", "Zamanında"), "PO-502": ("H2", "GECİKECEK"), "PO-503": ("H4", "GECİKECEK"), "PO-504": ("H1", "Zamanında"),
                          "PO-505": ("H2", "Riskli"), "PO-506": ("H4", "GECİKECEK"), "PO-507": ("", "Atanamadı")})
        self.assertEqual([x.no for x in self.s["siparisler"]][:2], ["PO-503", "PO-502"])
        self.assertEqual(self.p["PO-504"].baslangic, self.p["PO-501"].bitis)       # gün ortasında devralır

    def test_gecikme(self):
        x = self.p["PO-502"]
        self.assertEqual((x.bitis, x.hedef_bitis, x.bolluk), (date(2026, 11, 12), date(2026, 11, 7), -4))
        self.assertIn("300 adet/gün gerekir", next(u for u in self.u["Sevk gecikecek"] if u["siparis"] == "PO-502")["aciklama"])
        self.assertIn("18 iş günü bekleme", next(u for u in self.u["Hat bekleniyor"] if u["siparis"] == "PO-506")["aciklama"])
        self.assertEqual(self.u["Atanan hat yok"][0]["siparis"], "PO-507")
        self.assertIn("H3", self.u["Boş hat"][0]["aciklama"])

    def test_hat_ozeti(self):
        o = {x["hat"].ad: x for x in self.s["hat_ozet"]}
        self.assertEqual(o["H2"]["bosalma"], date(2026, 12, 1))
        self.assertEqual(o["H3"]["siparis"], 0)
        self.assertAlmostEqual(o["H1"]["yuklenen_dk"], 12000 * 12 + 4000 * 16)

    def test_excel(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Yükleme Planı", "Hat Yükü", "Günlük Plan", "Uyarılar"])
        gp = wb["Günlük Plan"]
        self.assertEqual(gp["B1"].value, "12.10")
        sutun = next(c.column for c in gp[1] if c.value == "29.10")
        self.assertEqual(gp.cell(3, sutun).value, "—")
        self.assertTrue(gp.cell(2, sutun - 1).value.endswith("½"))
        self.assertEqual(wb["Yükleme Planı"]["Q1"].value, "Planlayıcı Notu")


class KuralTesti(unittest.TestCase):
    def test_uzmanlik_disi_ve_ogrenme_yok(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t / "h.csv").write_text("Hat;Operatör Sayısı;Günlük Çalışma (dk);Verimlilik %;Ürün Grupları\nA;10;600;50;Elbise\nB;0;600;50;\n", encoding="utf-8")
            (t / "s.csv").write_text("Sipariş No;Ürün Grubu;Adet;SAM (dk);Sevk Tarihi;Atanan Hat\nX;Tişört;300;10;30.10.2026;A\nY;Tişört;100;10;30.10.2026;\n",
                                     encoding="utf-8")
            s = main.calistir(t / "h.csv", t / "s.csv", t / "o.xlsx", BAS, ogrenme=[], calisma_gunu=5)
            p = {x.no: x for x in s["siparisler"]}
            self.assertEqual((p["X"].baslangic, p["X"].bitis), (BAS, BAS))          # 3000 dk = 1 gün
            self.assertEqual(p["Y"].durum, "Atanamadı")
            turler = {u["tur"] for u in s["uyarilar"]}
            self.assertTrue({"Hat uzmanlığı dışı", "Uygun hat yok", "Hat kapasitesi sıfır", "Kesim tarihi yok"} <= turler)

    def test_takvim(self):
        t = main.Takvim(6)
        self.assertEqual(t.katsayi(date(2026, 5, 26)), 0.5)                 # Kurban Bayramı arifesi
        self.assertEqual(t.katsayi(date(2026, 5, 27)), 0)
        self.assertEqual(t.is_gunu_say(date(2026, 10, 26), date(2026, 10, 31)), 4)
        self.assertEqual(t.geri(date(2026, 11, 10), 2), date(2026, 11, 7))

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--hatlar", str(Path(tmp) / "yok.csv")]), 1)
            self.assertEqual(main.main(["--baslangic", "12/13/2026", "--cikti", str(Path(tmp) / "x.xlsx")]), 1)


if __name__ == "__main__":
    unittest.main()
