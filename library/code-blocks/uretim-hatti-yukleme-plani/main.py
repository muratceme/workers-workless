"""
Üretim Hattı Yükleme Planı — Workers / Workless kod bloğu
Tekstil ve Konfeksiyon › Planlama › Planlama Uzmanı

Siparişlerin dakika değeri (SAM) ve hat kapasitesinden hat bazında yükleme planı ve tahmini bitiş tarihleri çıkarır:
  - Hat günlük kapasitesi (dk) = operatör sayısı × günlük çalışma dakikası × verimlilik
  - Günlük adet = kapasite ÷ SAM; yeni modelin ilk günlerinde öğrenme eğrisi (varsayılan %50, %70, %85)
  - Siparişler sevk tarihine göre (en erken önce) sıralanır; her sipariş ürün grubuna uygun hatlar içinde en erken
    bitireceği hatta yüklenir. "Atanan Hat" yazılan sipariş o hatta kalır.
  - Sipariş, kesim hazır tarihinden ve hattın müsait olduğu tarihten önce başlamaz.
  - Pazar (isteğe göre Cumartesi), resmî ve dini bayramlar çalışılmaz; arife günleri yarım gün sayılır.
  - Tahmini bitiş + sevk öncesi tampon (varsayılan 2 iş günü: final kontrol, yükleme) sevk tarihiyle karşılaştırılır.
Rapor: yükleme planı, hat doluluğu, hat × gün adet planı, uyarılar. İnternete bağlanmaz.

Kullanım:
    python main.py                                                   # örnek: 4 hat, 7 sipariş, başlangıç 12.10.2026
    python main.py --hatlar hatlar.xlsx --siparisler siparisler.xlsx --baslangic 12.10.2026 --ogrenme 60,80 --tampon 3
"""
from __future__ import annotations

import argparse
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import puantaj_cekirdek as pc

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"

HAT_SUTUNLARI = {"hat": ("hat", "hat adi", "bant", "line"), "operator": ("operator sayisi", "operator", "kisi sayisi", "makineci sayisi"),
                 "dakika": ("gunluk calisma dk", "gunluk calisma", "gunluk dakika", "calisma dk"), "verim": ("verimlilik", "verim", "verimlilik yuzde"),
                 "gruplar": ("urun gruplari", "uzmanlik", "urun grubu"), "musait": ("musait oldugu tarih", "musait", "bosalma tarihi")}
SIPARIS_SUTUNLARI = {"no": ("siparis no", "po", "po no"), "model": ("model", "style"), "musteri": ("musteri", "buyer"),
                     "grup": ("urun grubu", "grup", "urun tipi"), "adet": ("adet", "miktar", "siparis adedi"),
                     "sam": ("sam dk", "sam", "dakika degeri", "birim sure dk"), "kesim": ("kesim hazir tarihi", "kesim hazir", "malzeme hazir tarihi", "baslangic"),
                     "sevk": ("sevk tarihi", "termin", "ex factory"), "atanan": ("atanan hat", "hat")}


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = pc.tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [pc.katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


def f(x, varsayilan=None) -> float | None:
    d = pc.sayi(x)
    return float(d) if d is not None else varsayilan


def oran(x, varsayilan=1.0) -> float:
    v = f(str(x).replace("%", "") if x not in (None, "") else None)
    if v is None:
        return varsayilan
    return v / 100 if v > 1 else v


def metin(x) -> str:
    return str(x if x is not None else "").strip()


# ----------------------------------------------------------------------------
# Takvim
# ----------------------------------------------------------------------------

class Takvim:
    def __init__(self, calisma_gunu: int = 6, ek_tatil: dict[date, str] | None = None):
        self.calisma_gunu = calisma_gunu            # 6: Pzt–Cmt, 5: Pzt–Cum
        self.ek = ek_tatil or {}
        self._yil: dict[int, tuple[dict, dict]] = {}

    def _t(self, d: date):
        if d.year not in self._yil:
            self._yil[d.year] = pc.tatiller(d.year, self.ek)
        return self._yil[d.year]

    def katsayi(self, d: date) -> float:
        """Günün çalışma katsayısı: 1 tam gün, 0,5 arife, 0 tatil / hafta sonu."""
        tam, yarim = self._t(d)
        if d.weekday() >= self.calisma_gunu or d in tam:
            return 0.0
        return 0.5 if d in yarim else 1.0

    def tatil_adi(self, d: date) -> str:
        tam, yarim = self._t(d)
        return tam.get(d) or yarim.get(d) or ""

    def ilk_is_gunu(self, d: date) -> date:
        while self.katsayi(d) == 0:
            d += timedelta(days=1)
        return d

    def geri(self, d: date, gun: int) -> date:
        while self.katsayi(d) == 0:
            d -= timedelta(days=1)
        while gun > 0:
            d -= timedelta(days=1)
            if self.katsayi(d) > 0:
                gun -= 1
        return d

    def is_gunu_say(self, a: date, b: date) -> int:
        """a (hariç) ile b (dahil) arasındaki iş günü; b < a ise negatif."""
        isaret, (x, y) = (1, (a, b)) if b >= a else (-1, (b, a))
        n, d = 0, x
        while d < y:
            d += timedelta(days=1)
            n += self.katsayi(d) > 0
        return isaret * n


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Hat:
    ad: str
    operator: float
    dakika: float
    verim: float
    gruplar: list
    musait: date | None
    # durum
    gun: date | None = None
    kullanilan: float = 0.0               # 'gun' içinde kullanılmış dakika
    toplam_dk: float = 0.0
    yuklenen: list = field(default_factory=list)

    @property
    def gunluk_dk(self) -> float:
        return self.operator * self.dakika * self.verim

    def uygun(self, grup: str) -> bool:
        return not self.gruplar or pc.katla(grup) in self.gruplar


@dataclass
class Siparis:
    satir: int
    no: str
    model: str
    musteri: str
    grup: str
    adet: int
    sam: float
    kesim: date | None
    sevk: date | None
    atanan: str
    hat: str = ""
    baslangic: date | None = None
    bitis: date | None = None
    gunluk: dict = field(default_factory=dict)        # tarih → adet
    hedef_bitis: date | None = None
    bolluk: int | None = None
    durum: str = ""
    gereken_gunluk: float | None = None


def hatlari_oku(yol: Path) -> dict[str, Hat]:
    sonuc = {}
    for r in kayitlar(yol, HAT_SUTUNLARI, ("hat", "operator", "dakika")):
        ad = metin(r.get("hat"))
        if not ad:
            continue
        gruplar = [pc.katla(g) for g in metin(r.get("gruplar")).replace(";", ",").split(",") if g.strip()]
        sonuc[ad] = Hat(ad, f(r.get("operator"), 0), f(r.get("dakika"), 0), oran(r.get("verim")), gruplar, pc.tarih(r.get("musait")))
    return sonuc


def siparisleri_oku(yol: Path) -> tuple[list[Siparis], list[str]]:
    sonuc, hatalar = [], []
    for r in kayitlar(yol, SIPARIS_SUTUNLARI, ("adet", "sam", "sevk")):
        no = metin(r.get("no")) or metin(r.get("model"))
        if not no:
            continue
        adet, sam = f(r.get("adet")), f(r.get("sam"))
        if not adet or not sam:
            hatalar.append(f"Satır {r['_satir']} ({no}): adet veya SAM okunamadı")
            continue
        sonuc.append(Siparis(r["_satir"], no, metin(r.get("model")), metin(r.get("musteri")), metin(r.get("grup")), int(adet), sam,
                             pc.tarih(r.get("kesim")), pc.tarih(r.get("sevk")), metin(r.get("atanan"))))
    return sonuc, hatalar


# ----------------------------------------------------------------------------
# Planlama
# ----------------------------------------------------------------------------

def simule(h: Hat, s: Siparis, baslangic: date, takvim: Takvim, ogrenme: list[float], gun_siniri: int = 730):
    """Siparişi hattın mevcut durumundan itibaren yükler. Dönüş: (başlangıç, bitiş, günlük adetler, son gün, son gün kullanılan dk)."""
    d, kullanilan = h.gun, h.kullanilan
    alt = max(x for x in (baslangic, h.musait, s.kesim) if x)
    if d is None or d < alt:
        d, kullanilan = alt, 0.0
    d = takvim.ilk_is_gunu(d)
    if kullanilan >= h.gunluk_dk * takvim.katsayi(d) - 1e-9:
        d, kullanilan = takvim.ilk_is_gunu(d + timedelta(days=1)), 0.0
    kalan = s.adet * s.sam                                  # üretilecek standart dakika
    gunluk, k, bas = {}, 0, None
    for _ in range(gun_siniri):
        kap = h.gunluk_dk * takvim.katsayi(d) - kullanilan
        if kap > 1e-9:
            bas = bas or d
            fakt = ogrenme[k] if k < len(ogrenme) else 1.0
            uretim = min(kalan, kap * fakt)
            gunluk[d] = gunluk.get(d, 0) + uretim / s.sam
            kalan -= uretim
            kullanilan += uretim / fakt
            k += 1
            if kalan <= 1e-6:
                return bas, d, gunluk, d, kullanilan
        d, kullanilan = takvim.ilk_is_gunu(d + timedelta(days=1)), 0.0
    raise ValueError(f"{s.no}: {gun_siniri} gün içinde bitirilemiyor; hat kapasitesini kontrol edin")


def planla(hatlar: dict[str, Hat], siparisler: list[Siparis], baslangic: date, takvim: Takvim, ogrenme: list[float], tampon: int = 2,
           risk_gun: int = 2) -> dict:
    uyarilar = []

    def uyar(onem, tur, s, aciklama):
        uyarilar.append({"onem": onem, "tur": tur, "siparis": s.no if s else "", "aciklama": aciklama})

    for h in hatlar.values():
        if h.gunluk_dk <= 0:
            uyar("Yüksek", "Hat kapasitesi sıfır", None, f"{h.ad}: operatör, çalışma dakikası veya verimlilik eksik; hat kullanılmadı")
    kullanilir = {k: h for k, h in hatlar.items() if h.gunluk_dk > 0}
    sirali = sorted(siparisler, key=lambda s: (s.sevk or date.max, s.kesim or date.min, s.no))
    for s in sirali:
        if s.atanan:
            if s.atanan not in kullanilir:
                s.durum = "Atanamadı"
                uyar("Yüksek", "Atanan hat yok", s, f"Atanan hat '{s.atanan}' hat listesinde yok veya kapasitesi sıfır; sipariş planlanmadı")
                continue
            adaylar = [kullanilir[s.atanan]]
            if not adaylar[0].uygun(s.grup):
                uyar("Orta", "Hat uzmanlığı dışı", s, f"{s.atanan} hattının ürün gruplarında '{s.grup}' yok; atama korunarak planlandı")
        else:
            adaylar = [h for h in kullanilir.values() if h.uygun(s.grup)]
            if not adaylar:
                s.durum = "Atanamadı"
                uyar("Yüksek", "Uygun hat yok", s, f"'{s.grup}' ürün grubunu dikebilen hat yok; sipariş planlanmadı")
                continue
        if not s.kesim:
            uyar("Bilgi", "Kesim tarihi yok", s, "Kesim hazır tarihi boş; plan başlangıcından itibaren yüklendi")
        secim = min(((simule(h, s, baslangic, takvim, ogrenme), h) for h in adaylar), key=lambda x: (x[0][1], x[0][4] / max(x[1].gunluk_dk, 1), x[1].ad))
        (bas, bit, gunluk, son_gun, son_kul), h = secim
        h.gun, h.kullanilan = son_gun, son_kul
        h.toplam_dk += s.adet * s.sam
        h.yuklenen.append(s)
        s.hat, s.baslangic, s.bitis, s.gunluk = h.ad, bas, bit, gunluk
        if s.kesim and bas > takvim.ilk_is_gunu(s.kesim) and (h.musait is None or bas > takvim.ilk_is_gunu(h.musait)):
            bekleme = takvim.is_gunu_say(takvim.ilk_is_gunu(max(s.kesim, baslangic)), bas)
            if bekleme > 0:
                uyar("Bilgi", "Hat bekleniyor", s, f"Kesim {s.kesim:%d.%m.%Y} hazır ama {h.ad} hattı {bas:%d.%m.%Y} tarihinde boşalıyor ({bekleme} iş günü bekleme)")
        if not s.sevk:
            s.durum = "Sevk tarihi yok"
            uyar("Orta", "Sevk tarihi yok", s, "Sevk tarihi boş; gecikme hesaplanamadı")
            continue
        s.hedef_bitis = takvim.geri(s.sevk, tampon)
        s.bolluk = takvim.is_gunu_say(bit, s.hedef_bitis)
        is_gunleri = sum(1 for g in gunluk)
        if s.bolluk < 0:
            s.durum = "GECİKECEK"
            gun_sayisi = takvim.is_gunu_say(bas - timedelta(days=1), s.hedef_bitis) if s.hedef_bitis >= bas else 0
            s.gereken_gunluk = s.adet / gun_sayisi if gun_sayisi > 0 else None
            uyar("Yüksek", "Sevk gecikecek", s, f"{h.ad}: tahmini bitiş {bit:%d.%m.%Y}, sevk {s.sevk:%d.%m.%Y} için en geç bitiş {s.hedef_bitis:%d.%m.%Y} "
                 f"({tampon} iş günü tampon) → {-s.bolluk} iş günü gecikme. Planlanan ortalama {s.adet / is_gunleri:.0f} adet/gün"
                 + (f"; zamanında bitmesi için {s.gereken_gunluk:.0f} adet/gün gerekir (ek hat, fazla mesai veya bölme)" if s.gereken_gunluk else
                    "; başlangıç zaten hedef bitişten sonra"))
        elif s.bolluk <= risk_gun:
            s.durum = "Riskli"
            uyar("Orta", "Az bolluk", s, f"{h.ad}: tahmini bitiş {bit:%d.%m.%Y}, en geç bitiş {s.hedef_bitis:%d.%m.%Y}; yalnız {s.bolluk} iş günü bolluk")
        else:
            s.durum = "Zamanında"
    son =max((s.bitis for s in siparisler if s.bitis), default=baslangic)
    hat_ozet = []
    for h in hatlar.values():
        bas = takvim.ilk_is_gunu(max(x for x in (baslangic, h.musait) if x))
        bitis = max((s.bitis for s in h.yuklenen), default=None)
        mevcut, d = 0.0, bas
        while d <= son:
            mevcut += h.gunluk_dk * takvim.katsayi(d)
            d += timedelta(days=1)
        hat_ozet.append({"hat": h, "siparis": len(h.yuklenen), "yuklenen_dk": h.toplam_dk, "mevcut_dk": mevcut,
                         "doluluk": h.toplam_dk / mevcut if mevcut else None, "bosalma": bitis, "baslangic": bas})
        if not h.yuklenen and h.gunluk_dk > 0:
            uyar("Bilgi", "Boş hat", None, f"{h.ad}: plan döneminde yüklenen sipariş yok ({bas:%d.%m.%Y} itibarıyla müsait)")
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    uyarilar.sort(key=lambda u: sira[u["onem"]])
    return {"siparisler": sirali, "hat_ozet": hat_ozet, "uyarilar": uyarilar, "son": son}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE", "GECİKECEK": "FDE2E1", "Atanamadı": "FDE2E1", "Riskli": "FFF4CE",
        "Zamanında": "E3F4E1"}
PALET = ["DCE9F9", "E3F4E1", "FCE8D5", "EFE3F7", "FFF4CE", "D9F2F2", "F7DCE4", "E8E8E8"]
KONTROL = PatternFill("solid", fgColor="FFF4CE")
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, basliklar, genislik=()):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def rapor_yaz(cikti: Path, s: dict, baslangic: date, takvim: Takvim, ogrenme: list[float], tampon: int) -> None:
    wb = Workbook()
    yp = wb.active
    yp.title = "Yükleme Planı"
    _baslik(yp, ["Sipariş No", "Model", "Müşteri", "Ürün Grubu", "Adet", "SAM", "Toplam dk", "Hat", "Kesim Hazır", "Başlangıç", "Tahmini Bitiş",
                 "En Geç Bitiş", "Sevk", "Bolluk (iş günü)", "Ort. Adet/Gün", "Durum", "Planlayıcı Notu"],
            (10, 8, 16, 11, 8, 6, 10, 6, 11, 11, 11, 11, 11, 9, 9, 12, 28))
    for x in s["siparisler"]:
        yp.append([x.no, x.model, x.musteri, x.grup, x.adet, x.sam, round(x.adet * x.sam), x.hat, x.kesim, x.baslangic, x.bitis, x.hedef_bitis,
                   x.sevk, x.bolluk, round(x.adet / len(x.gunluk)) if x.gunluk else None, x.durum, ""])
        r = yp.max_row
        for j in (9, 10, 11, 12, 13):
            yp.cell(r, j).number_format = "DD.MM.YYYY"
        if x.durum in RENK:
            yp.cell(r, 16).fill = PatternFill("solid", fgColor=RENK[x.durum])
        yp.cell(r, 17).fill = KONTROL
    yp.append([])
    yp.append([f"Plan başlangıcı {baslangic:%d.%m.%Y} · öğrenme eğrisi " + (", ".join(f"%{o * 100:.0f}" for o in ogrenme) or "yok")
               + f" · sevk öncesi tampon {tampon} iş günü · çalışma günü haftada {takvim.calisma_gunu}"])

    hy = wb.create_sheet("Hat Yükü")
    _baslik(hy, ["Hat", "Operatör", "Günlük Çalışma (dk)", "Verimlilik", "Günlük Kapasite (dk)", "Ürün Grupları", "Müsait", "Sipariş", "Yüklenen dk",
                 "Mevcut dk (plan sonuna kadar)", "Doluluk", "Boşalma Tarihi"], (6, 9, 10, 9, 11, 22, 11, 8, 11, 13, 9, 12))
    for x in s["hat_ozet"]:
        h = x["hat"]
        hy.append([h.ad, h.operator, h.dakika, h.verim, round(h.gunluk_dk), ", ".join(h.gruplar) or "Tümü", x["baslangic"], x["siparis"],
                   round(x["yuklenen_dk"]), round(x["mevcut_dk"]), x["doluluk"], x["bosalma"]])
        r = hy.max_row
        hy.cell(r, 4).number_format = hy.cell(r, 11).number_format = "0%"
        hy.cell(r, 7).number_format = hy.cell(r, 12).number_format = "DD.MM.YYYY"

    gp = wb.create_sheet("Günlük Plan")
    gunler, d = [], baslangic
    while d <= s["son"]:
        gunler.append(d)
        d += timedelta(days=1)
    gp.append(["Hat"] + [g.strftime("%d.%m") for g in gunler])
    gp.append(["Gün"] + [("Paz", "Pzt", "Sal", "Çar", "Per", "Cum", "Cmt")[(g.weekday() + 1) % 7] + (" ½" if takvim.katsayi(g) == 0.5 else "")
                         for g in gunler])
    for h in gp[1] + gp[2]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    renk = {x.no: PALET[i % len(PALET)] for i, x in enumerate(s["siparisler"])}
    for hs in s["hat_ozet"]:
        h = hs["hat"]
        satir = [h.ad]
        hucre_renk = {}
        for j, g in enumerate(gunler, 2):
            parcalar = [(x.no, x.gunluk[g]) for x in h.yuklenen if g in x.gunluk]
            if takvim.katsayi(g) == 0:
                satir.append("—")
            else:
                satir.append(" + ".join(f"{no}: {a:.0f}" for no, a in parcalar))
                if parcalar:
                    hucre_renk[j] = renk[parcalar[-1][0]]
        gp.append(satir)
        for j, c in hucre_renk.items():
            gp.cell(gp.max_row, j).fill = PatternFill("solid", fgColor=c)
    gp.column_dimensions["A"].width = 6
    for j in range(2, len(gunler) + 2):
        gp.column_dimensions[get_column_letter(j)].width = 14
    gp.freeze_panes = "B3"

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Sipariş", "Açıklama", "Karar"], (9, 22, 10, 100, 24))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["siparis"], u["aciklama"], ""])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).alignment = UST
        uy.cell(uy.max_row, 5).fill = KONTROL
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(hat_yolu: Path, siparis_yolu: Path, cikti: Path, baslangic: date | None = None, ogrenme: list[float] | None = None,
             tampon: int = 2, calisma_gunu: int = 6, ek_tatil: dict | None = None) -> dict:
    ogrenme = [0.5, 0.7, 0.85] if ogrenme is None else ogrenme
    takvim = Takvim(calisma_gunu, ek_tatil)
    hatlar = hatlari_oku(hat_yolu)
    siparisler, hatalar = siparisleri_oku(siparis_yolu)
    baslangic = takvim.ilk_is_gunu(baslangic or date.today())
    s = planla(hatlar, siparisler, baslangic, takvim, ogrenme, tampon)
    for h in hatalar:
        s["uyarilar"].append({"onem": "Orta", "tur": "Okunamayan satır", "siparis": "", "aciklama": h})
    rapor_yaz(cikti, s, baslangic, takvim, ogrenme, tampon)
    return {**s, "hatlar": hatlar, "baslangic": baslangic}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="SAM ve hat kapasitesinden hat bazında yükleme planı ve tahmini bitiş tarihleri hesaplar.")
    p.add_argument("--hatlar", type=Path, default=ORNEK / "hatlar.csv", help="Hat, Operatör Sayısı, Günlük Çalışma (dk), Verimlilik %%, Ürün Grupları, Müsait Olduğu Tarih")
    p.add_argument("--siparisler", type=Path, default=ORNEK / "siparisler.csv", help="Sipariş No, Model, Ürün Grubu, Adet, SAM, Kesim Hazır, Sevk, [Atanan Hat]")
    p.add_argument("--baslangic", help="Plan başlangıcı GG.AA.YYYY (varsayılan bugün; örnek veride 12.10.2026)")
    p.add_argument("--ogrenme", default="50,70,85", help="Yeni modelin ilk günlerinde verimlilik yüzdeleri, ör. 50,70,85 (kapatmak için boş: --ogrenme \"\")")
    p.add_argument("--tampon", type=int, default=2, help="Dikim bitişi ile sevk arası iş günü (final kontrol, yükleme; varsayılan 2)")
    p.add_argument("--calisma-gunu", type=int, choices=(5, 6), default=6, help="Haftalık çalışma günü (6: Pzt–Cmt, 5: Pzt–Cum)")
    p.add_argument("--ek-tatil", type=Path, help="İsteğe bağlı: Tarih, Açıklama (fabrika tatili, bakım duruşu)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "hat_yukleme_plani.xlsx")
    a = p.parse_args(argv)
    for y in (a.hatlar, a.siparisler, a.ek_tatil):
        if y is not None and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        ogrenme = [oran(x) for x in a.ogrenme.split(",") if x.strip()]
    except (TypeError, ValueError):
        print("[X] --ogrenme virgülle ayrılmış yüzdeler olmalı, ör. 50,70,85")
        return 1
    bas = pc.tarih(a.baslangic) if a.baslangic else (date(2026, 10, 12) if a.siparisler == ORNEK / "siparisler.csv" else None)
    if a.baslangic and bas is None:
        print("[X] --baslangic GG.AA.YYYY biçiminde olmalı")
        return 1
    try:
        s = calistir(a.hatlar, a.siparisler, a.cikti, bas, ogrenme, a.tampon, a.calisma_gunu, pc.ek_tatil_oku(a.ek_tatil) if a.ek_tatil else None)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(s['siparisler'])} sipariş · {len(s['hatlar'])} hat · plan {s['baslangic']:%d.%m.%Y} – {s['son']:%d.%m.%Y}")
    for x in s["siparisler"]:
        if x.bitis:
            print(f"[{'!' if x.durum in ('GECİKECEK', 'Riskli') else 'OK'}] {x.no} → {x.hat}: {x.baslangic:%d.%m} – {x.bitis:%d.%m.%Y} · {x.durum}"
                  + (f" ({-x.bolluk} iş günü)" if x.durum == "GECİKECEK" else ""))
        else:
            print(f"[X] {x.no}: {x.durum}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
