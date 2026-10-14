import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
GUN = date(2026, 9, 30)


def a(uye, ay, gun, tutar):
    return {"uye": uye, "ad": "", "tarih": date(2026, ay, gun), "tutar": D(tutar)}


def veri():
    return [a("U1", 1, 1, 400), a("U1", 2, 1, 300), a("u1", 3, 1, 300), a("U2", 6, 1, 100), a("U2", 7, 1, 100), a("U2", 8, 1, 100), a("U3", 7, 1, 700), a("U3", 7, 15, 700), a("U3", 7, 29, 600),
            a("U4", 5, 1, 100), a("U5", 9, 20, 50), a("U6", 9, 15, 200), a("U6", 9, 15, 100), a("U6", 9, 25, 200), a("U7", 10, 1, 999)]


UYELER = {"u1": {"ad": "Bir", "izin": "Evet", "kanal": "SMS"}, "u3": {"ad": "", "izin": "Hayır", "kanal": ""}, "u4": {"ad": "", "izin": "Evet", "kanal": ""}}


class AnalizTesti(unittest.TestCase):
    def test_durum_ve_liste(self):
        r = main.analiz(veri(), GUN, UYELER)
        self.assertEqual([(s["uye"], s["ad"], s["durum"], s["deger"], s["alisveris"], s["fis"], s["aralik"], s["once"], s["toplam"], s["sepet"], s["izin"]) for s in r["uyeler"]],
                         [("U3", "", "Risk altında", "Yüksek", 3, 3, D("14.0"), 63, D(2000), D("666.67"), "Hayır"), ("U1", "Bir", "Kayıp", "Yüksek", 3, 3, D("29.5"), 213, D(1000), D("333.33"), "Evet"),
                          ("U6", "", "Aktif", "Orta", 2, 3, D("10.0"), 5, D(500), D("166.67"), "Kayıt yok"), ("U2", "", "Aktif", "Düşük", 3, 3, D("30.5"), 60, D(300), D("100.00"), "Kayıt yok"),
                          ("U4", "", "Tek alışveriş", "Düşük", 1, 1, None, 152, D(100), D("100.00"), "Evet"), ("U5", "", "Aktif", "Düşük", 1, 1, None, 10, D(50), D("50.00"), "Kayıt yok")])
        self.assertEqual(([s["uye"] for s in r["hedef"]], [s["uye"] for s in r["liste"]]), (["U3", "U4", "U1"], ["U4", "U1"]))
        self.assertEqual(r["ozet"], {"Aktif": {"uye": 3, "toplam": D(850)}, "Risk altında": {"uye": 1, "toplam": D(2000)}, "Tek alışveriş": {"uye": 1, "toplam": D(100)}, "Kayıp": {"uye": 1, "toplam": D(1000)}})
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]), [("Analiz tarihinden sonraki alışveriş", ""), ("Yüksek değerli üye risk altında", "U3")])

    def test_esikler(self):
        d = lambda **k: {s["uye"]: s["durum"] for s in main.analiz(veri(), GUN, UYELER, **k)["uyeler"]}   # noqa: E731
        self.assertEqual(d(kat=D("1.5")), {"U3": "Risk altında", "U1": "Kayıp", "U6": "Aktif", "U2": "Risk altında", "U4": "Tek alışveriş", "U5": "Aktif"})
        self.assertEqual(d(kayip=213, tek=152, en_az=64), {"U3": "Aktif", "U1": "Risk altında", "U6": "Aktif", "U2": "Aktif", "U4": "Aktif", "U5": "Aktif"})
        r = main.analiz(veri(), GUN)
        self.assertEqual((r["liste"], {s["izin"] for s in r["uyeler"]}, "İzin bilgisi verilmedi" in [u["tur"] for u in r["uyarilar"]]), ([], {"Bilinmiyor"}, True))
        r = main.analiz(veri(), GUN, {"u1": UYELER["u1"]})
        self.assertEqual(sorted(u["kim"] for u in r["uyarilar"] if u["tur"] == "Üye kaydı yok"), ["U3", "U4"])
        with self.assertRaises(ValueError):
            main.analiz([a("X", 10, 5, 1)], GUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "uye_alisverisleri.csv", cikti, GUN, ORNEK / "uyeler.csv")
            self.assertEqual({d: o["uye"] for d, o in r["ozet"].items()}, {"Aktif": 9, "Risk altında": 5, "Tek alışveriş": 3, "Kayıp": 3})
            self.assertEqual([(s["uye"], s["durum"]) for s in r["liste"]], [("U1009", "Risk altında"), ("U1011", "Risk altında"), ("U1008", "Risk altında"), ("U1015", "Tek alışveriş"),
                                                                          ("U1016", "Tek alışveriş"), ("U1017", "Tek alışveriş"), ("U1014", "Kayıp"), ("U1013", "Kayıp"), ("U1012", "Kayıp")])
            self.assertEqual((len(r["hedef"]), r["uyarilar"]), (11, []))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Kampanya Listesi", "Tüm Hedef Üyeler", "Tüm Üyeler", "Özet", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, z = Path(tmp) / "a.csv", Path(tmp) / "u.csv"
            y.write_text("Kart No;Fiş Tarihi;Harcama\nK1;01.09.2026;1.250,50\nK1;x;5\nK2;01.09.2026;y\n;01.09.2026;5\n", encoding="utf-8")
            z.write_text("Kart No;İYS İzni\nK1;Onaylı\nK2;belki\nK3;\n", encoding="utf-8")
            l, uy = main.oku_alisveris(y)
            self.assertEqual(([(x["uye"], x["tarih"], x["tutar"]) for x in l], len(uy)), ([("K1", date(2026, 9, 1), D("1250.50"))], 2))
            d, uy = main.oku_uye(z)
            self.assertEqual(({k: v["izin"] for k, v in d.items()}, [u["tur"] for u in uy]), ({"k1": "Evet", "k2": "Hayır", "k3": "Hayır"}, ["İzin okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--alisverisler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--uyeler", str(z), "--tarih", "30.09.2026", "--kayip", "120", "--kat", "1,5", "--tek", "60", "--min", "20"]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--kat", "0"], ["--tek", "-1"])], [2] * 3)
            self.assertEqual((main.main(arg + ["--tarih", "01.01.2026"]), main.main(arg + ["--uyeler", str(Path(tmp) / "yok.csv")])), (1, 1))
            y.write_text("Kart No;Harcama\nK1;5\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
