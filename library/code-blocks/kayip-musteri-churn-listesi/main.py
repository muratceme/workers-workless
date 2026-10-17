"""
Kayıp Müşteri (Churn) Listesi — Gerisi Bende kod bloğu
CRM ve Sadakat Programı › CRM Uzmanı
Alışveriş sıklığı düşen sadakat programı üyelerini belirleyip geri kazanım kampanyası listesi çıkarır.
  Alışveriş sayısı     = farklı alışveriş günü sayısı (aynı gün birden çok fiş tek ziyaret sayılır)
  Ortalama aralık      = (son alışveriş − ilk alışveriş) ÷ (alışveriş sayısı − 1) gün
  Son alışverişten beri = analiz tarihi − son alışveriş (gün)
  Durum                = Kayıp: --kayip (180) günden uzun süredir alışveriş yok
                         Risk altında: en az iki alışverişi var ve son alışverişten beri geçen süre ortalama aralığın
                                       --kat (2) katından ve --min (30) günden uzun
                         Tek alışveriş: tek alışverişi var ve üzerinden --tek (90) günden fazla geçmiş
                         Aktif: diğerleri
  Değer segmenti       = toplam harcama sıralamasında ilk %20 Yüksek, sonraki %30 Orta, kalan Düşük
  Kampanya listesi     = Risk altında, Tek alışveriş ve Kayıp üyelerden ileti izni "Evet" olanlar
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

ALISVERIS = {"uye": ("uye no", "musteri no", "kart no", "uye", "musteri kodu"), "ad": ("ad soyad", "musteri adi", "uye adi", "ad"), "tarih": ("tarih", "alisveris tarihi", "fis tarihi"),
             "tutar": ("tutar", "ciro", "fis tutari", "harcama", "net tutar")}
UYE = {"uye": ("uye no", "musteri no", "kart no", "uye", "musteri kodu"), "ad": ("ad soyad", "musteri adi", "uye adi", "ad"), "izin": ("ileti izni", "izin", "ticari ileti izni", "iys izni", "pazarlama izni"),
       "kanal": ("kanal", "iletisim kanali", "izinli kanal")}
SIFIR = Decimal(0)
EVET, HAYIR = {"evet", "e", "var", "x", "1", "onay", "onayli", "izinli"}, {"hayir", "h", "yok", "0", "ret", "red", "izinsiz"}
DURUM_SIRA = {"Risk altında": 0, "Tek alışveriş": 1, "Kayıp": 2, "Aktif": 3}


def oku_alisveris(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, ALISVERIS, ("uye", "tarih", "tutar")):
        uye = metin(r["uye"])
        if not uye:
            continue
        t, tutar = tarih(r["tarih"]), para(r["tutar"])
        if t is None or tutar is None:
            uy.append(uyari("Yüksek", "Alışveriş satırı okunamadı", uye, f"Satır {r['_satir']}: tarih ya da tutar okunamadı. Satır alınmadı."))
            continue
        l.append({"uye": uye, "ad": metin(r["ad"]), "tarih": t, "tutar": tutar})
    return l, uy


def oku_uye(yol: Path) -> tuple[dict, list[dict]]:
    uy, d = [], {}
    for r in kayitlar(yol, UYE, ("uye", "izin")):
        uye = metin(r["uye"])
        if not uye:
            continue
        k = katla(r["izin"])
        izin = "Evet" if k in EVET else "Hayır" if k in HAYIR or not k else None
        if izin is None:
            uy.append(uyari("Orta", "İzin okunamadı", uye, f"Satır {r['_satir']}: ileti izni '{metin(r['izin'])}' anlaşılamadı; izin yok sayıldı."))
            izin = "Hayır"
        d[katla(uye)] = {"ad": metin(r["ad"]), "izin": izin, "kanal": metin(r["kanal"])}
    return d, uy


def analiz(alisverisler: list[dict], gun: date, uyeler: dict | None = None, kayip: int = 180, kat: Decimal = Decimal(2), tek: int = 90, en_az: int = 30) -> dict:
    gecerli = [a for a in alisverisler if a["tarih"] <= gun]
    uy = []
    if len(gecerli) < len(alisverisler):
        uy.append(uyari("Orta", "Analiz tarihinden sonraki alışveriş", "", f"{len(alisverisler) - len(gecerli)} alışveriş {gun:%d.%m.%Y} tarihinden sonra; hesaba alınmadı."))
    if not gecerli:
        raise ValueError("Analiz tarihine kadar alışveriş kaydı yok.")
    g = {}
    for a in gecerli:
        x = g.setdefault(katla(a["uye"]), {"uye": a["uye"], "ad": "", "gunler": set(), "toplam": SIFIR, "fis": 0, "son12": SIFIR})
        x["gunler"].add(a["tarih"])
        x["toplam"] += a["tutar"]
        x["fis"] += 1
        x["ad"] = a["ad"] or x["ad"]
        if (gun - a["tarih"]).days < 365:
            x["son12"] += a["tutar"]
    sonuc = []
    for k, x in g.items():
        gunler = sorted(x["gunler"])
        n, ilk, son = len(gunler), gunler[0], gunler[-1]
        aralik = yuvarla(Decimal((son - ilk).days) / (n - 1), "0.1") if n > 1 else None
        once = (gun - son).days
        if once > kayip:
            durum = "Kayıp"
        elif n > 1 and once > max(aralik * kat, Decimal(en_az)):
            durum = "Risk altında"
        elif n == 1 and once > tek:
            durum = "Tek alışveriş"
        else:
            durum = "Aktif"
        u = (uyeler or {}).get(k)
        sonuc.append({"uye": x["uye"], "ad": x["ad"] or (u["ad"] if u else ""), "alisveris": n, "fis": x["fis"], "ilk": ilk, "son": son, "aralik": aralik, "once": once, "toplam": x["toplam"],
                      "sepet": yuvarla(x["toplam"] / x["fis"]), "son12": x["son12"], "durum": durum, "izin": u["izin"] if u else ("Bilinmiyor" if uyeler is None else "Kayıt yok"),
                      "kanal": u["kanal"] if u else ""})
    sonuc.sort(key=lambda s: (-s["toplam"], katla(s["uye"])))
    yuksek, orta = math.ceil(len(sonuc) * 0.2), math.ceil(len(sonuc) * 0.5)
    for i, s in enumerate(sonuc, 1):
        s["deger"] = "Yüksek" if i <= yuksek else "Orta" if i <= orta else "Düşük"
    hedef = sorted((s for s in sonuc if s["durum"] != "Aktif"), key=lambda s: (DURUM_SIRA[s["durum"]], -s["toplam"], katla(s["uye"])))
    liste = [s for s in hedef if s["izin"] == "Evet"]
    if uyeler is None:
        uy.append(uyari("Orta", "İzin bilgisi verilmedi", "", "Üye dosyası (--uyeler) verilmediği için kampanya listesi boş. Ticari elektronik ileti yalnız onay alınmış üyelere gönderilebilir."))
    else:
        for s in hedef:
            if s["izin"] == "Kayıt yok":
                uy.append(uyari("Bilgi", "Üye kaydı yok", s["uye"], "Alışverişi olan üye, üye dosyasında yok; izin bilinmediği için listeye alınmadı."))
    ozet = {d: {"uye": sum(1 for s in sonuc if s["durum"] == d), "toplam": sum((s["toplam"] for s in sonuc if s["durum"] == d), SIFIR)} for d in ("Aktif", "Risk altında", "Tek alışveriş", "Kayıp")}
    for s in hedef:
        if s["deger"] == "Yüksek" and s["durum"] == "Risk altında":
            uy.append(uyari("Orta", "Yüksek değerli üye risk altında", s["uye"], f"{s['once']} gündür alışveriş yok (ortalama aralık {s['aralik']} gün); toplam harcama {s['toplam']:,.2f}."))
    return {"uyeler": sonuc, "hedef": hedef, "liste": liste, "ozet": ozet, "uyarilar": uy}


DURUM_RENK = {"Risk altında": SARI, "Tek alışveriş": MAVI, "Kayıp": KIRMIZI, "Aktif": YESIL}
BASLIK = ["Üye No", "Ad Soyad", "Durum", "Değer", "Alışveriş Günü", "Fiş", "İlk Alışveriş", "Son Alışveriş", "Ortalama Aralık (gün)", "Son Alışverişten Beri (gün)", "Toplam Harcama",
          "Ortalama Fiş", "Son 12 Ay Harcama", "İleti İzni", "Kanal"]
GENISLIK = (13, 22, 13, 8, 9, 6, 11, 11, 10, 11, 14, 11, 14, 10, 12)


def _yaz(ws, s):
    _satir(ws, [s["uye"], s["ad"], s["durum"], s["deger"], s["alisveris"], s["fis"], s["ilk"], s["son"], s["aralik"], s["once"], s["toplam"], s["sepet"], s["son12"], s["izin"], s["kanal"]],
           para_s=(11, 12, 13), tarih_s=(7, 8), dolgu={3: DURUM_RENK[s["durum"]], 14: None if s["izin"] == "Evet" else SARI})


def rapor_yaz(cikti: Path, r: dict, gun: date) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kampanya Listesi"
    _baslik(ws, BASLIK, GENISLIK, "C2")
    for s in r["liste"]:
        _yaz(ws, s)
    ws.append([])
    ws.append([f"Analiz tarihi {gun:%d.%m.%Y}. Yalnız ileti izni 'Evet' olan Risk altında, Tek alışveriş ve Kayıp üyeler. Gönderimden önce İYS kaydını doğrulayın."])
    ws = wb.create_sheet("Tüm Hedef Üyeler")
    _baslik(ws, BASLIK, GENISLIK, "C2")
    for s in r["hedef"]:
        _yaz(ws, s)
    ws = wb.create_sheet("Tüm Üyeler")
    _baslik(ws, BASLIK, GENISLIK, "C2")
    for s in r["uyeler"]:
        _yaz(ws, s)
    ws = wb.create_sheet("Özet")
    _baslik(ws, ["Durum", "Üye Sayısı", "Toplam Harcama"], (16, 10, 16))
    for d, o in r["ozet"].items():
        _satir(ws, [d, o["uye"], o["toplam"]], para_s=(3,), dolgu={1: DURUM_RENK[d]})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(alisveris_yol: Path, cikti: Path, gun: date, uye_yol: Path | None = None, kayip: int = 180, kat: Decimal = Decimal(2), tek: int = 90, en_az: int = 30) -> dict:
    alisverisler, uy = oku_alisveris(alisveris_yol)
    uyeler, uy2 = oku_uye(uye_yol) if uye_yol else (None, [])
    r = analiz(alisverisler, gun, uyeler, kayip, kat, tek, en_az)
    r["uyarilar"] = sirala(uy + uy2 + r["uyarilar"])
    rapor_yaz(cikti, r, gun)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Alışveriş sıklığı düşen sadakat programı üyelerini belirleyip geri kazanım kampanyası listesi çıkarır.")
    p.add_argument("--alisverisler", type=Path, help="Üye No, Ad Soyad, Tarih, Tutar")
    p.add_argument("--uyeler", type=Path, help="Üye No, Ad Soyad, İleti İzni (Evet / Hayır), Kanal")
    p.add_argument("--tarih", help="Analiz tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 30.09.2026)")
    p.add_argument("--kayip", type=int, default=180, help="Bu kadar günden uzun süredir alışveriş yapmayan üye kayıp (varsayılan 180)")
    p.add_argument("--kat", default="2", help="Son alışverişten beri geçen süre ortalama aralığın bu katını aşarsa risk (varsayılan 2)")
    p.add_argument("--tek", type=int, default=90, help="Tek alışverişli üye bu kadar gün dönmezse listelenir (varsayılan 90)")
    p.add_argument("--min", type=int, default=30, help="Risk için en az geçen gün (varsayılan 30)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kayip_musteri_listesi.xlsx")
    a = p.parse_args(argv)
    ornek = a.alisverisler is None
    ay = a.alisverisler or ORNEK / "uye_alisverisleri.csv"
    uy = a.uyeler or (ORNEK / "uyeler.csv" if ornek else None)
    gun = tarih(a.tarih) if a.tarih else (date(2026, 9, 30) if ornek else date.today())
    kat = para(a.kat)
    if gun is None or kat is None or kat <= 0 or min(a.kayip, a.tek, a.min) < 0:
        print("[X] --tarih GG.AA.YYYY, --kat sıfırdan büyük, gün sayıları sıfır ya da üstü olmalı")
        return 2
    for y in (ay, uy):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(ay, a.cikti, gun, uy, a.kayip, kat, a.tek, a.min)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {len(r['uyeler'])} üye · aktif {o['Aktif']['uye']} · risk altında {o['Risk altında']['uye']} · tek alışveriş {o['Tek alışveriş']['uye']} · kayıp {o['Kayıp']['uye']}")
    print(f"[OK] Kampanya listesi: {len(r['liste'])} üye (izinli) · hedef üyelerden {len(r['hedef']) - len(r['liste'])} kişi izinsiz ya da izni bilinmiyor")
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
