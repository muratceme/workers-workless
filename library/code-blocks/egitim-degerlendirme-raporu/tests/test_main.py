import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def k(egitim, kisi, katildi=True, on=None, son=None, mem=None, dep="D"):
    f = lambda v: None if v is None else D(str(v))   # noqa: E731
    return {"egitim": egitim, "kisi": kisi, "departman": dep, "katildi": katildi, "on": f(on), "son": f(son), "memnuniyet": f(mem)}


A = lambda: [k("A", "K1", True, 40, 80, 5, "Finans"), k("A", "K2", True, 60, 70, 4, "Finans"), k("A", "K3", True, 50, 65, 3, "Satış"), k("A", "K4", True, None, 90, None, "Satış"), k("A", "K5", False, dep="Satış")]   # noqa: E731
B = lambda: [k("B", "K3", True, 70, 60, 2, "Satış"), k("B", "K6", True, 80, 80, 3, "Pazarlama")]   # noqa: E731
AL = ("davetli", "katilan", "katilim", "on", "son", "gelisim", "cift", "sinava_giren", "gecen", "basari", "memnuniyet", "anket")


class AnalizTesti(unittest.TestCase):
    def test_egitim_ozeti(self):
        r = main.analiz(A() + B())
        self.assertEqual([(o["egitim"], tuple(o[a] for a in AL), o["isaret"]) for o in r["egitimler"]],
                         [("B", (2, 2, D("100.00"), D("75.00"), D("70.00"), D("-5.00"), 2, 2, 1, D("50.00"), D("2.50"), 2), ["Başarı düşük", "Memnuniyet düşük", "Gelişim yok"]),
                          ("A", (5, 4, D("80.00"), D("50.00"), D("76.25"), D("21.67"), 3, 4, 3, D("75.00"), D("4.00"), 3), [])])
        self.assertEqual(tuple(r["genel"][a] for a in AL), (7, 6, D("85.71"), D("60.00"), D("74.17"), D("11.00"), 5, 6, 4, D("66.67"), D("3.40"), 5))
        self.assertEqual(([(u["tur"], u["kim"]) for u in r["uyarilar"]], r["katilmayan"], [(x["egitim"], x["kisi"], x["son"]) for x in r["basarisiz"]]),
                         ([("Eğitim eşiklerin altında", "B")], [{"egitim": "A", "kisi": "K5", "departman": "Satış"}], [("A", "K3", D(65)), ("B", "K3", D(60))]))

    def test_departman_ve_esikler(self):
        r = main.analiz(A() + B())
        self.assertEqual([(o["departman"], o["davetli"], o["katilan"], o["katilim"], o["son"], o["basari"], o["memnuniyet"]) for o in r["departmanlar"]],
                         [("Finans", 2, 2, D("100.00"), D("75.00"), D("100.00"), D("4.50")), ("Pazarlama", 1, 1, D("100.00"), D("80.00"), D("100.00"), D("3.00")), ("Satış", 4, 3, D("75.00"), D("71.67"), D("33.33"), D("2.50"))])
        i = lambda **kw: main.analiz(A(), **kw)["egitimler"][0]["isaret"]   # noqa: E731
        self.assertEqual((i(), i(katilim_esik=D(81)), i(basari_esik=D(76)), i(memnuniyet_esik=D("4.1")), i(gecme=D(81)), i(gecme=D(81), basari_esik=D(25))), ([], ["Katılım düşük"], ["Başarı düşük"], ["Memnuniyet düşük"], ["Başarı düşük"], []))
        r = main.analiz([k("C", "X"), k("C", "Y", False), k("C", "Z", False), k("D", "X", True, 50, 50), k("E", "X", True, None, None, 4), k("E", "Y"), k("E", "Z")])
        self.assertEqual([(o["egitim"], o["katilim"], o["on"], o["basari"], o["memnuniyet"], o["gelisim"], o["isaret"]) for o in r["egitimler"]],
                         [("C", D("33.33"), None, None, None, None, ["Katılım düşük", "Ölçüm yok"]), ("D", D("100.00"), D("50.00"), D("0.00"), None, D("0.00"), ["Başarı düşük", "Gelişim yok"]), ("E", D("100.00"), None, None, D("4.00"), None, [])])
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]), [("Anket yanıt oranı düşük", "E"), ("Eğitim eşiklerin altında", "C"), ("Eğitim eşiklerin altında", "D")])
        with self.assertRaises(ValueError):
            main.analiz([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "e.xlsx"
            r = main.calistir(ORNEK / "egitim_katilim.csv", cikti)
            self.assertEqual([(o["egitim"], o["katilim"], o["gelisim"], o["basari"], o["memnuniyet"], o["isaret"]) for o in r["egitimler"]],
                             [("Etkili Sunum", D("100.00"), D("-5.00"), D("50.00"), D("2.50"), ["Başarı düşük", "Memnuniyet düşük", "Gelişim yok"]), ("Zaman Yönetimi", D("33.33"), None, None, None, ["Katılım düşük", "Ölçüm yok"]),
                              ("İleri Excel", D("80.00"), D("21.67"), D("75.00"), D("4.00"), [])])
            self.assertEqual((r["genel"]["davetli"], r["genel"]["katilan"], len(r["katilmayan"]), len(r["basarisiz"]), len(r["uyarilar"])), (10, 7, 3, 2, 2))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Eğitim Özeti", "Departman Özeti", "Takip Listesi", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "k.csv"
            y.write_text("Konu;Personel;Katılım;Ön Test;Sınav;Anket\nA;P1;;45,5;72.5;4\nA;P2;;;;\nA;P3;Hayır;;80;\nA;P4;belki;;;\nA;P5;Evet;;120;\nA;P6;Evet;;;6\nA;P7;katıldı;x;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["kisi"], x["katildi"], x["on"], x["son"], x["memnuniyet"]) for x in l], sorted(u["tur"] for u in uy)),
                             ([("P1", True, D("45.5"), D("72.5"), D(4)), ("P2", False, None, None, None), ("P3", False, None, None, None)], ["Katılmadı ama puanı var"] + ["Satır okunamadı"] * 4))
            self.assertEqual([x["kisi"] for x in main.oku(y, D(10))[0]], ["P1", "P2", "P3", "P6"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--katilim", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--gecme", "60", "--olcek", "10", "--katilim-esik", "50", "--basari-esik", "50", "--memnuniyet-esik", "7,5"]), 0)
            for ek in (["--gecme", "101"], ["--olcek", "0"], ["--memnuniyet-esik", "6"], ["--katilim-esik", "x"]):
                self.assertEqual(main.main(arg + ek), 2)
            self.assertEqual(main.main(["--katilim", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Eğitim;Puan\nA;1\n", "Eğitim;Katılımcı;Katıldı\nA;;Evet\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
