"""
Eğitim Değerlendirme Raporu — Gerisi Bende kod bloğu
İnsan Kaynakları › Eğitim ve Gelişim Uzmanı

Katılım, ön test, son test ve memnuniyet anketi verisinden eğitim bazında etkinlik raporu üretir.
  Katılım %          = katılan ÷ davet edilen × 100
  Ön / son test ort. = puanı girilmiş katılımcıların ortalaması
  Gelişim (puan)     = hem ön hem son testi olan katılımcıların (son test ort. − ön test ort.)
  Başarı %           = son testi --gecme (70) ve üzeri olan ÷ son testi olan × 100
  Memnuniyet         = anket puanlarının ortalaması (--olcek 5 üzerinden) ve yanıt sayısı
  İşaretler          = katılım < --katilim-esik (%80) · başarı < --basari-esik (%70) · memnuniyet < --memnuniyet-esik (3,5) · gelişim ≤ 0
Rapor tepki (memnuniyet) ve öğrenme (test) düzeyini ölçer; eğitimin iş başındaki davranışa ve sonuçlara etkisini göstermez. Eşikler varsayımdır, kurumunuza göre değiştirin.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"egitim": ("egitim", "egitim adi", "konu"), "kisi": ("katilimci", "ad soyad", "personel", "calisan"), "departman": ("departman", "birim"), "katildi": ("katildi", "katilim", "katilim durumu"),
         "on": ("on test", "on test puani", "on degerlendirme"), "son": ("son test", "son test puani", "sinav", "son degerlendirme"), "memnuniyet": ("memnuniyet", "anket", "memnuniyet puani", "anket puani")}
EVET, HAYIR = {"evet", "e", "x", "katildi", "var", "1", "true"}, {"hayir", "h", "katilmadi", "yok", "0", "false", "mazeretli", "izinli"}
Z = Decimal(0)


def olcu(x) -> Decimal | None:
    """Puan: 72,5 = 72.5."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    try:
        return Decimal(re.sub(r"[^\d,.\-]", "", str(x)).replace(",", "."))
    except InvalidOperation:
        return None


def oku(yol: Path, olcek: Decimal = Decimal(5)) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("egitim", "kisi")):
        e, k = metin(r["egitim"]), metin(r["kisi"])
        if not e and not k:
            continue
        on, son, mem = (olcu(r[a]) if metin(r[a]) else None for a in ("on", "son", "memnuniyet"))
        d = katla(r["katildi"])
        bozuk = [ad for ad, alan, v, ust in (("ön test", "on", on, 100), ("son test", "son", son, 100), ("memnuniyet", "memnuniyet", mem, olcek)) if (v is None and metin(r[alan])) or (v is not None and not 0 <= v <= ust)]
        if not e or not k or bozuk or (d and d not in EVET | HAYIR):
            uy.append(uyari("Yüksek", "Satır okunamadı", f"{e or '?'} · {k or '?'} (satır {r['_satir']})", f"Katıldı '{metin(r['katildi'])}' Evet veya Hayır olmalı; test puanları 0–100, memnuniyet 0–{olcek.normalize():f} arasında olmalı"
                            + (f" (sorunlu: {', '.join(bozuk)})" if bozuk else "") + ". Satır alınmadı."))
            continue
        veri = any(v is not None for v in (on, son, mem))
        katildi = d in EVET if d else veri
        if d in HAYIR and veri:
            uy.append(uyari("Orta", "Katılmadı ama puanı var", f"{e} · {k}", "Katılım 'Hayır' yazılmış ama test veya anket puanı girilmiş. Katılmamış sayıldı, puanlar hesaba katılmadı; kaydı düzeltin."))
            on = son = mem = None
        l.append({"egitim": e, "kisi": k, "departman": metin(r["departman"]) or "Belirtilmemiş", "katildi": katildi, "on": on, "son": son, "memnuniyet": mem})
    return l, uy


def ort(d: list[Decimal]) -> Decimal | None:
    return yuvarla(sum(d, Z) / len(d)) if d else None


def analiz(kayit: list[dict], gecme: Decimal = Decimal(70), katilim_esik: Decimal = Decimal(80), basari_esik: Decimal = Decimal(70), memnuniyet_esik: Decimal = Decimal("3.5")) -> dict:
    if not kayit:
        raise ValueError("Değerlendirilecek kayıt yok.")
    uy = []

    def ozet(grup: list[dict]) -> dict:
        kat = [x for x in grup if x["katildi"]]
        cift = [x for x in kat if x["on"] is not None and x["son"] is not None]
        sonlu = [x["son"] for x in kat if x["son"] is not None]
        mem = [x["memnuniyet"] for x in kat if x["memnuniyet"] is not None]
        gelisim = yuvarla(sum((x["son"] - x["on"] for x in cift), Z) / len(cift)) if cift else None
        return {"davetli": len(grup), "katilan": len(kat), "katilim": yuvarla(Decimal(len(kat)) / len(grup) * 100), "on": ort([x["on"] for x in kat if x["on"] is not None]), "son": ort(sonlu), "gelisim": gelisim, "cift": len(cift),
                "gecen": sum(1 for s in sonlu if s >= gecme), "sinava_giren": len(sonlu), "basari": yuvarla(Decimal(sum(1 for s in sonlu if s >= gecme)) / len(sonlu) * 100) if sonlu else None, "memnuniyet": ort(mem), "anket": len(mem)}

    egitimler, eg = [], {}
    for x in kayit:
        eg.setdefault(katla(x["egitim"]), {"ad": x["egitim"], "kayit": []})["kayit"].append(x)
    for g in eg.values():
        o = {"egitim": g["ad"]} | ozet(g["kayit"])
        isaret = []
        if o["katilim"] < katilim_esik:
            isaret.append("Katılım düşük")
        if o["basari"] is not None and o["basari"] < basari_esik:
            isaret.append("Başarı düşük")
        if o["memnuniyet"] is not None and o["memnuniyet"] < memnuniyet_esik:
            isaret.append("Memnuniyet düşük")
        if o["gelisim"] is not None and o["gelisim"] <= 0:
            isaret.append("Gelişim yok")
        if o["katilan"] and not o["sinava_giren"] and not o["anket"]:
            isaret.append("Ölçüm yok")
        o["isaret"] = isaret
        egitimler.append(o)
        if isaret:
            ayrinti = f"katılım %{o['katilim']}" + (f", başarı %{o['basari']}" if o["basari"] is not None else "") + (f", memnuniyet {o['memnuniyet']}" if o["memnuniyet"] is not None else "") + (f", gelişim {o['gelisim']} puan" if o["gelisim"] is not None else "")
            uy.append(uyari("Orta", "Eğitim eşiklerin altında", g["ad"], f"{', '.join(isaret)} ({ayrinti}). İçeriği, eğitmeni, süreyi ve hedef kitleyi gözden geçirin."))
        if o["anket"] and o["anket"] < o["katilan"] / 2:
            uy.append(uyari("Bilgi", "Anket yanıt oranı düşük", g["ad"], f"{o['katilan']} katılımcıdan {o['anket']} tanesi anketi doldurmuş; memnuniyet ortalaması temsil edici olmayabilir."))
    dep = {}
    for x in kayit:
        dep.setdefault(katla(x["departman"]), {"ad": x["departman"], "kayit": []})["kayit"].append(x)
    departmanlar = sorted(({"departman": g["ad"]} | ozet(g["kayit"]) for g in dep.values()), key=lambda o: o["departman"])
    katilmayan = [{"egitim": x["egitim"], "kisi": x["kisi"], "departman": x["departman"]} for x in kayit if not x["katildi"]]
    basarisiz = [{"egitim": x["egitim"], "kisi": x["kisi"], "departman": x["departman"], "son": x["son"]} for x in kayit if x["katildi"] and x["son"] is not None and x["son"] < gecme]
    egitimler.sort(key=lambda o: (not o["isaret"], o["egitim"]))
    return {"egitimler": egitimler, "departmanlar": departmanlar, "katilmayan": katilmayan, "basarisiz": basarisiz, "genel": {"egitim": "Genel"} | ozet(kayit), "ayar": (gecme, katilim_esik, basari_esik, memnuniyet_esik), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    gecme, ke, be, me = r["ayar"]
    ws = wb.active
    ws.title = "Eğitim Özeti"
    ust = ["Davetli", "Katılan", "Katılım %", "Ön Test Ort.", "Son Test Ort.", "Gelişim (puan)", "Sınava Giren", "Geçen", "Başarı %", "Memnuniyet", "Anket Yanıtı"]
    deger = lambda o: [o["davetli"], o["katilan"], o["katilim"], o["on"], o["son"], o["gelisim"], o["sinava_giren"] or None, o["gecen"] if o["sinava_giren"] else None, o["basari"], o["memnuniyet"], o["anket"] or None]   # noqa: E731
    _baslik(ws, ["Eğitim"] + ust + ["İşaretler", "Aksiyon (siz doldurun)"], (36, 8, 8, 10, 11, 11, 11, 9, 8, 10, 11, 9, 40, 30), "B2")
    for o in r["egitimler"] + [r["genel"]]:
        n = _satir(ws, [o["egitim"]] + deger(o) + [", ".join(o.get("isaret", [])), ""], para_s=(4, 5, 6, 7, 10, 11),
                   dolgu={4: KIRMIZI if o["katilim"] < ke else None, 10: KIRMIZI if o["basari"] is not None and o["basari"] < be else None, 11: KIRMIZI if o["memnuniyet"] is not None and o["memnuniyet"] < me else None,
                          7: KIRMIZI if o["gelisim"] is not None and o["gelisim"] <= 0 else None}, sar=(13,))
        if o is r["genel"]:
            for c in ws[n]:
                c.font = Font(bold=True)
    ws.append([])
    ws.append([f"Geçme puanı {gecme.normalize():f}. Eşikler: katılım %{ke.normalize():f}, başarı %{be.normalize():f}, memnuniyet {me.normalize():f}. Gelişim yalnız hem ön hem son testi olan katılımcılardan hesaplanır."])
    ws = wb.create_sheet("Departman Özeti")
    _baslik(ws, ["Departman"] + ust, (26, 8, 8, 10, 11, 11, 11, 9, 8, 10, 11, 9), "B2")
    for o in r["departmanlar"]:
        _satir(ws, [o["departman"]] + deger(o), para_s=(4, 5, 6, 7, 10, 11))
    ws = wb.create_sheet("Takip Listesi")
    _baslik(ws, ["Durum", "Eğitim", "Katılımcı", "Departman", "Son Test", "Not (siz doldurun)"], (16, 36, 28, 22, 10, 30))
    for x in r["katilmayan"]:
        _satir(ws, ["Katılmadı", x["egitim"], x["kisi"], x["departman"], None, ""], dolgu={1: SARI})
    for x in r["basarisiz"]:
        _satir(ws, ["Geçme puanı altı", x["egitim"], x["kisi"], x["departman"], x["son"], ""], dolgu={1: KIRMIZI})
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, olcek: Decimal = Decimal(5), **ayar) -> dict:
    kayit, uy = oku(yol, olcek)
    r = analiz(kayit, **ayar)
    uy += r["uyarilar"]
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Katılım, sınav ve memnuniyet verisinden eğitim etkinlik raporu üretir.")
    p.add_argument("--katilim", type=Path, default=ORNEK / "egitim_katilim.csv", help="Eğitim, Katılımcı, Departman, Katıldı, Ön Test, Son Test, Memnuniyet")
    p.add_argument("--gecme", default="70", help="Son test geçme puanı (varsayılan 70)")
    p.add_argument("--olcek", default="5", help="Memnuniyet anketinin en yüksek puanı (varsayılan 5)")
    p.add_argument("--katilim-esik", default="80", help="Katılım bu %%'nin altındaysa işaretlenir (varsayılan 80)")
    p.add_argument("--basari-esik", default="70", help="Başarı oranı bu %%'nin altındaysa işaretlenir (varsayılan 70)")
    p.add_argument("--memnuniyet-esik", default="3,5", help="Memnuniyet ortalaması bunun altındaysa işaretlenir (varsayılan 3,5)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "egitim_degerlendirme_raporu.xlsx")
    a = p.parse_args(argv)
    g, o, ke, be, me = (olcu(x) for x in (a.gecme, a.olcek, a.katilim_esik, a.basari_esik, a.memnuniyet_esik))
    if None in (g, o, ke, be, me) or not 0 <= g <= 100 or o <= 0 or not 0 <= ke <= 100 or not 0 <= be <= 100 or not 0 <= me <= o:
        print("[X] --gecme, --katilim-esik ve --basari-esik 0–100; --olcek pozitif; --memnuniyet-esik 0 ile ölçek arasında olmalı")
        return 2
    if not a.katilim.exists():
        print(f"[X] Dosya bulunamadı: {a.katilim}")
        return 1
    try:
        r = calistir(a.katilim, a.cikti, o, gecme=g, katilim_esik=ke, basari_esik=be, memnuniyet_esik=me)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    gn = r["genel"]
    print(f"[OK] {len(r['egitimler'])} eğitim · {gn['davetli']} davet · katılım %{gn['katilim']}" + (f" · başarı %{gn['basari']}" if gn["basari"] is not None else "") + (f" · memnuniyet {gn['memnuniyet']}" if gn["memnuniyet"] is not None else ""))
    for e in r["egitimler"]:
        if e["isaret"]:
            print(f"[!] {e['egitim']}: {', '.join(e['isaret'])}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
