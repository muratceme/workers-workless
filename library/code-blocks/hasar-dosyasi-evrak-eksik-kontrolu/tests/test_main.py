import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 8)


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "h.xlsx"
        cls.s = main.calistir(ORNEK / "dosyalar.csv", ORNEK / "evraklar.csv", None, main.VARSAYILAN_LISTE, cls.cikti, BUGUN)
        cls.d = {x.no: x for x in cls.s["dosyalar"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def eksik(self, no, ic=False):
        return [k.evrak for k in self.d[no].eksik if k.ic == ic]

    def test_durumlar(self):
        self.assertEqual({n: x.durum for n, x in self.d.items()}, {
            "H-2026-0101": "Eksik evrak (3) · Hatırlatma", "H-2026-0102": "Eksik evrak (3) · Hatırlatma", "H-2026-0103": "Tamam",
            "H-2026-0104": "Eksik evrak (1) · Hatırlatma", "H-2026-0105": "Eksik evrak (3) · Hatırlatma",
            "H-2026-0106": "Eksik evrak (3) · Hatırlatma", "H-2026-0107": "İç evrak bekleniyor", "H-2026-0108": "Eksik evrak (2)",
            "H-2026-0109": "Evrak listesi tanımsız"})

    def test_kosullu_evraklar(self):
        self.assertIn("Alkol raporu", self.eksik("H-2026-0101"))                       # resmi tutanak
        self.assertIn("Rehin kaldırma (fek) yazısı veya rehin alacaklı muvafakatnamesi", self.eksik("H-2026-0102"))
        self.assertEqual(self.eksik("H-2026-0104"), ["Vekaletname"])
        self.assertEqual(self.eksik("H-2026-0105"), ["Banka hesap bilgisi (IBAN)", "Maluliyet (sağlık kurulu) raporu", "Gelir belgesi"])
        self.assertNotIn("Ölüm belgesi", [k.evrak for k in self.d["H-2026-0105"].gerekli])
        self.assertIn("Bina yönetimi / komşu tutanağı", [k.evrak for k in self.d["H-2026-0108"].gerekli])

    def test_eslestirme(self):
        e = {x.ad: x.karsiladigi for x in self.d["H-2026-0101"].evraklar}
        self.assertEqual(e["Ehliyet ve ruhsat fotokopisi"], ["Sürücü belgesi", "Araç ruhsatı"])
        self.assertEqual(e["Trafik kazası tespit tutanağı (polis)"], ["Kaza tespit tutanağı"])
        self.assertEqual(e["Servis randevu belgesi"], [])
        self.assertEqual(self.eksik("H-2026-0107", ic=True), ["Ekspertiz raporu"])
        self.assertTrue(main.kelime_eslesir("kaza tespit", "Kazası tespit edildi"))
        self.assertFalse(main.kelime_eslesir("gelir", "Geliş belgesi"))

    def test_cikti(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Özet", "Dosyalar", "Eksik Evraklar", "Gelen Evraklar", "Talep Yazıları"])
        self.assertEqual(wb["Talep Yazıları"].max_row, 7)                            # 6 eksikli dosya + başlık
        yazi = (self.s["yazilar"] / "H-2026-0101.txt").read_text(encoding="utf-8")
        self.assertIn("Sayın Ali Kurgu", yazi)
        self.assertIn("- Alkol raporu", yazi)
        self.assertNotIn("Ekspertiz", yazi)                                            # iç takip, sigortalıdan istenmez
        self.assertFalse((self.s["yazilar"] / "H-2026-0103.txt").exists())

    def test_hatirlatma_gun(self):
        s = main.calistir(ORNEK / "dosyalar.csv", ORNEK / "evraklar.csv", None, main.VARSAYILAN_LISTE, self.cikti, BUGUN, hatirlatma_gun=30)
        self.assertEqual(sum("Hatırlatma" in x.durum for x in s["dosyalar"]), 1)     # yalnız 0102 (35 gün)


class KlasorTesti(unittest.TestCase):
    def test_klasor_ve_uyarilar(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t / "dosyalar.csv").write_text("Dosya No;Branş;Hasar Türü;Sigortalı;Hasar Tarihi;İhbar Tarihi\nK1;Kasko;Cam kırılması;Deneme;05.10.2026;01.10.2026\n",
                                            encoding="utf-8")
            for alt, adlar in {"K1": ["hasar_ihbar_formu.pdf", "ruhsat.jpg", "cam-faturasi.pdf"], "K9": ["x.pdf"]}.items():
                (t / "evrak" / alt).mkdir(parents=True)
                for a in adlar:
                    (t / "evrak" / alt / a).write_bytes(b"x")
            s = main.calistir(t / "dosyalar.csv", None, t / "evrak", main.VARSAYILAN_LISTE, t / "o" / "o.xlsx", BUGUN)
            d = s["dosyalar"][0]
            self.assertEqual([k.evrak for k in d.eksik if not k.ic], ["Banka hesap bilgisi (IBAN)", "Hasar fotoğrafları"])
            self.assertIn("K9: evrakı var ama dosya listesinde yok", s["uyarilar"])
            self.assertIn("K1: ihbar tarihi hasar tarihinden önce", s["uyarilar"])

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--tarih", "yarın", "--cikti", str(Path(tmp) / "x.xlsx")]), 2)


if __name__ == "__main__":
    unittest.main()
