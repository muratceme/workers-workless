"""
Hasar Dosyası Evrak Eksik Kontrolü — Workers / Workless kod bloğu
Sigorta › Hasar Yönetimi › Hasar Uzman Yardımcısı

Branş ve hasar türüne göre gerekli evrak listesini (evrak_listesi.csv; şirketinize göre düzenleyin) dosyaya gelen
evraklarla karşılaştırır:
  - Gerekli evraklar dosyanın branşına, hasar türüne ve koşullarına (vekil, tüzel kişi, rehinli, resmi tutanak,
    yaralanma, ölüm, komşu…) göre seçilir.
  - Gelen evrak adları anahtar kelimelerle eşleştirilir ("Ehliyet ve ruhsat fotokopisi" iki evrakı karşılar);
    tanınmayan evraklar ayrıca listelenir.
  - Eksikler sigortalıdan / hak sahibinden istenecekler ve iç takip (ör. ekspertiz raporu) olarak ayrılır.
  - İhbardan bu yana geçen gün; eksik evrakı olan ve --hatirlatma-gun'ü geçen dosyalar "Hatırlatma" alır.
  - Her eksikli dosya için evrak talep yazısı taslağı (.txt) üretilir.
Gelen evraklar bir tablodan (Dosya No, Evrak, Geliş Tarihi) veya dosya numarası adlı alt klasörlerdeki dosya
adlarından okunur. İnternete bağlanmaz.

Kullanım:
    python main.py                                                     # örnek dosyalar (rapor tarihi 08.10.2026)
    python main.py --dosyalar dosyalar.xlsx --evraklar evraklar.xlsx --tarih 09.10.2026
    python main.py --dosyalar dosyalar.xlsx --klasor ./hasar_dosyalari --evrak-listesi sirket_evrak_listesi.csv
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
VARSAYILAN_LISTE = BURASI / "evrak_listesi.csv"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")

DOSYA_SUTUNLARI = {
    "no": ("dosya no", "hasar dosya no", "hasar no", "dosya"),
    "brans": ("brans", "urun", "police turu"),
    "tur": ("hasar turu", "hasar tipi", "tur"),
    "sigortali": ("sigortali", "hak sahibi", "sigortali adi", "musteri"),
    "hasar_tarihi": ("hasar tarihi", "olay tarihi"),
    "ihbar_tarihi": ("ihbar tarihi", "acilis tarihi", "basvuru tarihi"),
    "kosullar": ("kosullar", "kosul", "ozel durum", "notlar"),
}
EVRAK_SUTUNLARI = {"no": DOSYA_SUTUNLARI["no"], "evrak": ("evrak", "evrak adi", "belge", "belge adi", "dosya adi"),
                   "tarih": ("gelis tarihi", "tarih", "teslim tarihi")}
LISTE_SUTUNLARI = {"brans": ("brans",), "tur": ("hasar turu",), "evrak": ("evrak",), "anahtar": ("anahtar kelimeler", "anahtar"),
                   "kosul": ("kosul",), "kimden": ("kimden",), "not": ("not", "aciklama")}
IC_KAYNAK = ("sirket", "eksper", "ic ")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9*]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(str(x or "").strip(), f).date()
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            break
    else:
        raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")
    return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} for r in s[bi + 1:]]


def coklu(x) -> set[str]:
    return {katla(p) for p in re.split(r"[|,;]", str(x or "")) if katla(p)}


def kelime_eslesir(anahtar: str, ad: str) -> bool:
    """Anahtarın kelimeleri addaki ardışık kelimelerin başında geçiyor mu? ('tutanak' → 'tutanağı', 'kaza' → 'kazası')."""
    kt, dt = katla(anahtar).split(), katla(ad).split()
    kokler = [k[:-1] if len(k) >= 6 else k for k in kt]
    return any(all(dt[i + j].startswith(k) for j, k in enumerate(kokler)) for i in range(len(dt) - len(kt) + 1))


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Kural:
    branslar: set[str]
    turler: set[str]
    evrak: str
    anahtarlar: list[str]
    kosullar: set[str]
    kimden: str
    not_: str

    @property
    def ic(self) -> bool:
        k = katla(self.kimden)
        return any(k.startswith(x.strip()) for x in IC_KAYNAK)

    def uygun(self, d: "Dosya") -> bool:
        b, t = katla(d.brans), katla(d.tur)
        return (("*" in self.branslar or b in self.branslar) and ("*" in self.turler or t in self.turler)
                and (not self.kosullar or bool(self.kosullar & d.kosullar)))

    def karsilar(self, evrak_adi: str) -> bool:
        return any(kelime_eslesir(a, evrak_adi) for a in self.anahtarlar)


@dataclass
class Evrak:
    ad: str
    tarih: date | None
    karsiladigi: list[str] = field(default_factory=list)


@dataclass
class Dosya:
    no: str
    brans: str
    tur: str
    sigortali: str
    hasar_tarihi: date | None
    ihbar_tarihi: date | None
    kosullar: set[str]
    evraklar: list[Evrak] = field(default_factory=list)
    gerekli: list[Kural] = field(default_factory=list)
    eksik: list[Kural] = field(default_factory=list)
    tanimli: bool = True
    gun: int | None = None
    durum: str = ""


def liste_oku(yol: Path) -> list[Kural]:
    return [Kural(coklu(k["brans"]) or {"*"}, coklu(k["tur"]) or {"*"}, str(k["evrak"]).strip(),
                  [a for a in str(k.get("anahtar") or k["evrak"]).split("|") if a.strip()], coklu(k.get("kosul")),
                  str(k.get("kimden") or "Sigortalı").strip(), str(k.get("not") or "").strip())
            for k in kayitlar(yol, LISTE_SUTUNLARI, ("brans", "tur", "evrak")) if k.get("evrak")]


def dosyalari_oku(yol: Path) -> list[Dosya]:
    return [Dosya(str(k["no"]).strip(), str(k.get("brans") or "").strip(), str(k.get("tur") or "").strip(), str(k.get("sigortali") or "").strip(),
                  tarih(k.get("hasar_tarihi")), tarih(k.get("ihbar_tarihi")), coklu(k.get("kosullar")))
            for k in kayitlar(yol, DOSYA_SUTUNLARI, ("no", "brans")) if k.get("no")]


def evraklari_oku(yol: Path | None, klasor: Path | None) -> dict[str, list[Evrak]]:
    sonuc = defaultdict(list)
    if yol:
        for k in kayitlar(yol, EVRAK_SUTUNLARI, ("no", "evrak")):
            if k.get("no") and k.get("evrak"):
                sonuc[str(k["no"]).strip()].append(Evrak(str(k["evrak"]).strip(), tarih(k.get("tarih"))))
    if klasor:
        for alt in sorted(p for p in klasor.iterdir() if p.is_dir()):
            for f in sorted(p for p in alt.rglob("*") if p.is_file()):
                sonuc[alt.name].append(Evrak(re.sub(r"[_\-]+", " ", f.stem), date.fromtimestamp(f.stat().st_mtime)))
    return sonuc


# ----------------------------------------------------------------------------
# Kontrol
# ----------------------------------------------------------------------------

def kontrol_et(dosyalar: list[Dosya], evraklar: dict[str, list[Evrak]], kurallar: list[Kural], bugun: date, hatirlatma_gun: int) -> list[str]:
    uyarilar = []
    bilinen = {d.no for d in dosyalar}
    for no in sorted(set(evraklar) - bilinen):
        uyarilar.append(f"{no}: evrakı var ama dosya listesinde yok")
    for d in dosyalar:
        d.evraklar = evraklar.get(d.no, [])
        d.gerekli = [k for k in kurallar if k.uygun(d)]
        d.tanimli = any("*" not in k.branslar for k in d.gerekli)
        for e in d.evraklar:
            e.karsiladigi = [k.evrak for k in d.gerekli if k.karsilar(e.ad)]
        d.eksik = [k for k in d.gerekli if not any(k.evrak in e.karsiladigi for e in d.evraklar)]
        d.gun = (bugun - d.ihbar_tarihi).days if d.ihbar_tarihi else None
        dis = [k for k in d.eksik if not k.ic]
        if not d.tanimli:
            d.durum = "Evrak listesi tanımsız"
        elif dis:
            d.durum = f"Eksik evrak ({len(dis)})" + (" · Hatırlatma" if d.gun is not None and d.gun >= hatirlatma_gun else "")
        elif d.eksik:
            d.durum = "İç evrak bekleniyor"
        else:
            d.durum = "Tamam"
        if d.hasar_tarihi and d.ihbar_tarihi and d.ihbar_tarihi < d.hasar_tarihi:
            uyarilar.append(f"{d.no}: ihbar tarihi hasar tarihinden önce")
    return uyarilar


def talep_yazisi(d: Dosya, bugun: date) -> str:
    dis = [k for k in d.eksik if not k.ic]
    satir = [f"Tarih: {bugun:%d.%m.%Y}", f"Konu: {d.no} numaralı hasar dosyası — eksik belgeler", "", f"Sayın {d.sigortali or 'Sigortalımız'},", "",
             (f"{d.hasar_tarihi:%d.%m.%Y} tarihli " if d.hasar_tarihi else "") + f"{d.brans} {d.tur.lower()} hasarınıza ilişkin {d.no} numaralı "
             "dosyanın değerlendirilebilmesi için aşağıdaki belgelerin tarafımıza iletilmesini rica ederiz:", ""]
    satir += [f"- {k.evrak}" + (f" ({k.not_})" if k.not_ else "") for k in dis]
    satir += ["", "Belgeler tamamlandığında dosyanızın değerlendirmesi sürdürülecektir. Sorularınız için hasar birimimizle iletişime "
                  "geçebilirsiniz.", "", "Saygılarımızla,", "[Şirket adı] Hasar Birimi"]
    return "\n".join(satir)


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
TAKIP = PatternFill("solid", fgColor="FFF4CE")
DURUM_RENK = {"Tamam": "E3F4E1", "İç evrak bekleniyor": "E8F0FE", "Evrak listesi tanımsız": "EEEEEE"}
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def durum_rengi(d: str) -> str:
    return DURUM_RENK.get(d, "FDE2E1" if "Hatırlatma" in d else "FFF4CE")


def rapor_yaz(cikti: Path, dosyalar: list[Dosya], uyarilar: list[str], bugun: date, liste_yolu: Path) -> Path:
    wb = Workbook()
    o = wb.active
    o.title = "Özet"
    o.column_dimensions["A"].width, o.column_dimensions["B"].width = 34, 70
    for a, b in [("Rapor tarihi", f"{bugun:%d.%m.%Y}"), ("Dosya sayısı", len(dosyalar)), ("Evrak listesi", liste_yolu.name),
                 ("Sigortalıdan istenecek evrak", sum(1 for d in dosyalar for k in d.eksik if not k.ic)),
                 ("İç takipteki evrak", sum(1 for d in dosyalar for k in d.eksik if k.ic)),
                 ("Tanınmayan gelen evrak", sum(1 for d in dosyalar for e in d.evraklar if not e.karsiladigi))]:
        o.append([a, b])
        o.cell(o.max_row, 1).font = Font(bold=True)
    o.append([])
    o.append(["Durum", "Dosya"])
    for h in o[o.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for durum, n in sorted(Counter(re.sub(r" \(\d+\)", "", d.durum) for d in dosyalar).items()):
        o.append([durum, n])
    if uyarilar:
        o.append([])
        for u in uyarilar:
            o.append(["Uyarı", u])
    o.append([])
    o.append(["Not", "Evrak listesi örnektir; şirketinizin hasar evrak listesine, genel şartlara ve dosyanın özelliğine göre güncelleyin."])
    o.cell(o.max_row, 2).alignment = UST

    ds = wb.create_sheet("Dosyalar")
    _baslik(ds, ["Dosya No", "Branş", "Hasar Türü", "Sigortalı / Hak Sahibi", "Hasar Tarihi", "İhbar Tarihi", "İhbardan Gün", "Koşullar",
                 "Gerekli", "Gelen", "Eksik (Sigortalıdan)", "Eksik (İç)", "Tanınmayan Evrak", "Durum", "Takip Notu"],
            (13, 9, 14, 24, 11, 11, 8, 18, 8, 7, 48, 22, 26, 22, 26))
    for d in dosyalar:
        ds.append([d.no, d.brans, d.tur, d.sigortali, d.hasar_tarihi, d.ihbar_tarihi, d.gun, ", ".join(sorted(d.kosullar)), len(d.gerekli), len(d.evraklar),
                   "\n".join(k.evrak for k in d.eksik if not k.ic) or "—", "\n".join(k.evrak for k in d.eksik if k.ic) or "—",
                   "\n".join(e.ad for e in d.evraklar if not e.karsiladigi) or "—", d.durum, ""])
        for c in (5, 6):
            ds.cell(ds.max_row, c).number_format = "DD.MM.YYYY"
        ds.cell(ds.max_row, 14).fill = PatternFill("solid", fgColor=durum_rengi(d.durum))
        ds.cell(ds.max_row, 15).fill = TAKIP
        for h in ds[ds.max_row]:
            h.alignment = UST
    ds.auto_filter.ref = f"A1:O{ds.max_row}"

    ek = wb.create_sheet("Eksik Evraklar")
    _baslik(ek, ["Dosya No", "Sigortalı / Hak Sahibi", "Evrak", "Kimden", "Not", "İhbardan Gün", "İstendi mi? / Tarih"], (13, 24, 44, 22, 44, 8, 18))
    for d in dosyalar:
        for k in d.eksik:
            ek.append([d.no, d.sigortali, k.evrak, k.kimden, k.not_, d.gun, ""])
            ek.cell(ek.max_row, 7).fill = TAKIP
    ek.auto_filter.ref = f"A1:G{ek.max_row}"

    ge = wb.create_sheet("Gelen Evraklar")
    _baslik(ge, ["Dosya No", "Evrak", "Geliş Tarihi", "Karşıladığı Gerekli Evrak"], (13, 44, 12, 60))
    for d in dosyalar:
        for e in d.evraklar:
            ge.append([d.no, e.ad, e.tarih, "; ".join(e.karsiladigi) or "TANINMADI — elle kontrol edin"])
            ge.cell(ge.max_row, 3).number_format = "DD.MM.YYYY"
            if not e.karsiladigi:
                ge.cell(ge.max_row, 4).fill = TAKIP

    ty = wb.create_sheet("Talep Yazıları")
    _baslik(ty, ["Dosya No", "Talep Yazısı Taslağı"], (13, 110))
    klasor = cikti.parent / "talep_yazilari"
    klasor.mkdir(parents=True, exist_ok=True)
    for d in dosyalar:
        if d.tanimli and any(not k.ic for k in d.eksik):
            metin = talep_yazisi(d, bugun)
            ty.append([d.no, metin])
            ty.cell(ty.max_row, 2).alignment = UST
            (klasor / f"{re.sub(r'[^A-Za-z0-9_-]+', '_', d.no)}.txt").write_text(metin + "\n", encoding="utf-8")
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    return klasor


def calistir(dosya_yolu: Path, evrak_yolu: Path | None, klasor: Path | None, liste_yolu: Path, cikti: Path, bugun: date | None = None,
             hatirlatma_gun: int = 15) -> dict:
    bugun = bugun or date.today()
    kurallar = liste_oku(liste_yolu)
    dosyalar = dosyalari_oku(dosya_yolu)
    uyarilar = kontrol_et(dosyalar, evraklari_oku(evrak_yolu, klasor), kurallar, bugun, hatirlatma_gun)
    yazilar = rapor_yaz(cikti, dosyalar, uyarilar, bugun, liste_yolu)
    return {"dosyalar": dosyalar, "uyarilar": uyarilar, "kurallar": kurallar, "yazilar": yazilar, "bugun": bugun}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Hasar dosyalarındaki eksik evrakları branş ve hasar türüne göre bulur, talep yazısı taslağı hazırlar.")
    p.add_argument("--dosyalar", type=Path, default=ORNEK / "dosyalar.csv", help="Hasar dosyaları (.xlsx/.csv)")
    p.add_argument("--evraklar", type=Path, help="Gelen evraklar: Dosya No, Evrak, Geliş Tarihi (.xlsx/.csv)")
    p.add_argument("--klasor", type=Path, help="Dosya numarası adlı alt klasörler; içindeki dosya adları gelen evrak sayılır")
    p.add_argument("--evrak-listesi", type=Path, default=VARSAYILAN_LISTE, help="Gerekli evrak listesi (varsayılan: evrak_listesi.csv, örnektir)")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY (varsayılan: bugün; örnek veride 08.10.2026)")
    p.add_argument("--hatirlatma-gun", type=int, default=15, help="İhbardan bu kadar gün sonra eksik evrak için 'Hatırlatma' (varsayılan 15)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "hasar_evrak_kontrolu.xlsx")
    a = p.parse_args(argv)
    ornek = a.dosyalar == ORNEK / "dosyalar.csv"
    evraklar = a.evraklar or (ORNEK / "evraklar.csv" if ornek and not a.klasor else None)
    bugun = tarih(a.tarih) if a.tarih else (date(2026, 10, 8) if ornek else None)
    if a.tarih and bugun is None:
        print(f"[X] Tarih GG.AA.YYYY olmalı: {a.tarih}")
        return 2
    for y in (a.dosyalar, evraklar, a.evrak_listesi):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    if not evraklar and not a.klasor:
        print("[X] Gelen evraklar için --evraklar veya --klasor verin.")
        return 2
    try:
        s = calistir(a.dosyalar, evraklar, a.klasor, a.evrak_listesi, a.cikti, bugun, a.hatirlatma_gun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    d = Counter(re.sub(r" \(\d+\)", "", x.durum) for x in s["dosyalar"])
    print(f"[OK] {len(s['dosyalar'])} dosya · " + " · ".join(f"{k} {v}" for k, v in sorted(d.items())))
    for x in s["dosyalar"]:
        if "Hatırlatma" in x.durum:
            print(f"[!] {x.no}: {x.gun} gündür eksik — " + ", ".join(k.evrak for k in x.eksik if not k.ic))
    for u in s["uyarilar"]:
        print(f"[!] {u}")
    print(f"[OK] Rapor: {a.cikti.resolve()} · talep yazıları: {s['yazilar'].resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
