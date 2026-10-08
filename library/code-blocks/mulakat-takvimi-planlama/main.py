"""
Mülakat Takvimi Planlama — Workers / Workless kod bloğu
İnsan Kaynakları › İnsan Kaynakları Uzman Yardımcısı

Aday ve mülakatçı müsaitliklerini eşleştirip çakışmasız mülakat takvimi ve davet listesi çıkarır:
  - Her pozisyon için gereken roller (ör. İK + Teknik) aynı anda panel olarak görüşür; süre ve format pozisyondan.
  - Mülakatçı yalnız yetkili olduğu pozisyonlara atanır; günlük en çok mülakat sayısı ve iki mülakat arası tampon
    (varsayılan 15 dk) korunur.
  - Sıralama: müsaitliği en dar aday önce; adayın en erken uygun saati seçilir (15 dk adımlarla); aynı roldeki
    mülakatçılardan en az yüklü olan atanır.
  - Yerleşemeyen adaylar için hangi rolün ortak müsaitliği olmadığı yazılır.
Rapor: takvim, mülakatçı programı ve yükü, yerleşemeyenler, davet metinleri; .ics takvim dosyası. İnternete
bağlanmaz.

Kullanım:
    python main.py                                                   # örnek: 9 aday, 5 mülakatçı, 3 pozisyon
    python main.py --adaylar adaylar.xlsx --mulakatcilar mulakatcilar.xlsx --pozisyonlar pozisyonlar.csv --tampon 10
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
ADIM = timedelta(minutes=15)
GUNLER = ("Pazartesi", "Salı", "Çarşamba", "Perşembe", "Cuma", "Cumartesi", "Pazar")
ARALIK_RE = re.compile(r"(\d{1,2}[./]\d{1,2}[./]\d{4}|\d{4}-\d{2}-\d{2})\s+(\d{1,2})[:.](\d{2})\s*[-–]\s*(\d{1,2})[:.](\d{2})")

ADAY_SUTUNLARI = {"no": ("aday no", "no"), "ad": ("ad soyad", "aday", "ad"), "pozisyon": ("pozisyon", "basvurulan pozisyon"),
                  "eposta": ("e posta", "eposta", "email"), "musaitlik": ("musaitlik", "uygun zamanlar")}
MULAKATCI_SUTUNLARI = {"ad": ("mulakatci", "ad soyad", "ad"), "rol": ("rol",), "pozisyonlar": ("pozisyonlar", "pozisyon"),
                       "gunluk": ("gunluk en cok", "gunluk maks", "gunluk en fazla"), "musaitlik": ("musaitlik", "uygun zamanlar")}
POZISYON_SUTUNLARI = {"ad": ("pozisyon",), "roller": ("gerekli roller", "roller", "panel"), "sure": ("sure dk", "sure"), "format": ("format", "mulakat formati")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def araliklar(x) -> list[tuple[datetime, datetime]]:
    sonuc = []
    for t, h1, m1, h2, m2 in ARALIK_RE.findall(str(x or "")):
        g = None
        for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d"):
            try:
                g = datetime.strptime(t, f)
                break
            except ValueError:
                pass
        if g:
            bas, bit = g.replace(hour=int(h1), minute=int(m1)), g.replace(hour=int(h2), minute=int(m2))
            if bit > bas:
                sonuc.append((bas, bit))
    return sorted(sonuc)


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=";"))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Son sütun (müsaitlik) içinde ';' olabileceği için fazladan hücreler son eşleşen sütuna birleştirilir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            son = len(b) - 1
            sonuc = []
            for n, r2 in enumerate(s[bi + 1:], bi + 2):
                r2 = list(r2)
                if len(r2) > len(b):
                    r2 = r2[:son] + ["; ".join(str(c) for c in r2[son:] if c not in (None, ""))]
                sonuc.append({a: (r2[j] if j < len(r2) else None) for a, j in es.items()} | {"_satir": n})
            return sonuc
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Pozisyon:
    ad: str
    roller: list[str]
    sure: int
    format: str


@dataclass
class Mulakatci:
    ad: str
    rol: str
    pozisyonlar: set[str]           # katlanmış; boş veya "tumu" = hepsi
    gunluk: int | None
    musait: list
    dolu: list = field(default_factory=list)       # (başlangıç, bitiş, aday)

    def yetkili(self, pozisyon: str) -> bool:
        return not self.pozisyonlar or "tumu" in self.pozisyonlar or katla(pozisyon) in self.pozisyonlar

    def bos_mu(self, bas: datetime, bit: datetime, tampon: timedelta) -> bool:
        if not any(a <= bas and bit <= b for a, b in self.musait):
            return False
        if self.gunluk and sum(1 for x in self.dolu if x[0].date() == bas.date()) >= self.gunluk:
            return False
        return all(bit + tampon <= x[0] or x[1] + tampon <= bas for x in self.dolu)


@dataclass
class Aday:
    no: str
    ad: str
    pozisyon: str
    eposta: str
    musait: list
    baslangic: datetime | None = None
    bitis: datetime | None = None
    panel: list = field(default_factory=list)
    neden: str = ""

    @property
    def musait_dk(self) -> float:
        return sum((b - a).total_seconds() / 60 for a, b in self.musait)


def oku_pozisyonlar(yol: Path) -> dict[str, Pozisyon]:
    return {katla(r["ad"]): Pozisyon(metin(r["ad"]), [x.strip() for x in re.split(r"[,+/]", metin(r.get("roller"))) if x.strip()] or ["İK"],
                                     int(float(metin(r.get("sure")) or 45)), metin(r.get("format")) or "Yüz yüze")
            for r in kayitlar(yol, POZISYON_SUTUNLARI, ("ad",)) if r.get("ad")}


def oku_mulakatcilar(yol: Path) -> list[Mulakatci]:
    sonuc = []
    for r in kayitlar(yol, MULAKATCI_SUTUNLARI, ("ad", "rol", "musaitlik")):
        if r.get("ad"):
            poz = {katla(x) for x in re.split(r"[,;]", metin(r.get("pozisyonlar"))) if x.strip()}
            g = metin(r.get("gunluk"))
            sonuc.append(Mulakatci(metin(r["ad"]), metin(r["rol"]), poz, int(float(g)) if g else None, araliklar(r.get("musaitlik"))))
    return sonuc


def oku_adaylar(yol: Path) -> list[Aday]:
    return [Aday(metin(r.get("no")) or f"Satır {r['_satir']}", metin(r.get("ad")), metin(r.get("pozisyon")), metin(r.get("eposta")),
                 araliklar(r.get("musaitlik")))
            for r in kayitlar(yol, ADAY_SUTUNLARI, ("ad", "pozisyon", "musaitlik")) if r.get("ad")]


# ----------------------------------------------------------------------------
# Planlama
# ----------------------------------------------------------------------------

def planla(adaylar: list[Aday], mulakatcilar: list[Mulakatci], pozisyonlar: dict[str, Pozisyon], tampon_dk: int = 15) -> dict:
    tampon = timedelta(minutes=tampon_dk)
    uyarilar = []
    for a in sorted(adaylar, key=lambda a: (a.musait_dk, a.no)):
        poz = pozisyonlar.get(katla(a.pozisyon))
        if poz is None:
            a.neden = f"'{a.pozisyon}' pozisyon tablosunda yok"
            continue
        if not a.musait:
            a.neden = "Müsaitlik okunamadı (biçim: GG.AA.YYYY SS:DD-SS:DD)"
            continue
        sure = timedelta(minutes=poz.sure)
        adaylar_rol = {rol: [m for m in mulakatcilar if katla(m.rol) == katla(rol) and m.yetkili(poz.ad)] for rol in poz.roller}
        eksik_rol = [rol for rol, lst in adaylar_rol.items() if not lst]
        if eksik_rol:
            a.neden = f"{', '.join(eksik_rol)} rolünde bu pozisyona yetkili mülakatçı yok"
            continue
        bulundu = None
        bos_rol = set()
        for bas_aralik, bit_aralik in a.musait:
            t = bas_aralik
            while t + sure <= bit_aralik and bulundu is None:
                panel = []
                for rol, lst in adaylar_rol.items():
                    uygun = [m for m in lst if m.bos_mu(t, t + sure, tampon)]
                    if uygun:
                        bos_rol.add(rol)
                        panel.append(min(uygun, key=lambda m: (len(m.dolu), m.ad)))
                if len(panel) == len(adaylar_rol):
                    bulundu = (t, panel)
                t += ADIM
            if bulundu:
                break
        if bulundu is None:
            gunler = ", ".join(f"{g:%d.%m.%Y}" for g in sorted({x[0].date() for x in a.musait}))
            hic = [rol for rol in poz.roller if rol not in bos_rol]
            a.neden = (f"{', '.join(hic)} rolünde adayın zamanlarında ({gunler}) boş {poz.sure} dk yok" if hic
                       else f"Rollerin ({', '.join(poz.roller)}) adayın zamanlarında ({gunler}) aynı anda boş olduğu {poz.sure} dk yok")
            continue
        t, panel = bulundu
        a.baslangic, a.bitis, a.panel = t, t + sure, panel
        for m in panel:
            m.dolu.append((t, t + sure, a))
            m.dolu.sort(key=lambda x: x[0])
    for m in mulakatcilar:
        if not m.musait:
            uyarilar.append(f"{m.ad}: müsaitlik okunamadı")
    yerlesen = sorted((a for a in adaylar if a.baslangic), key=lambda a: (a.baslangic, a.no))
    return {"yerlesen": yerlesen, "yerlesemeyen": [a for a in adaylar if not a.baslangic], "uyarilar": uyarilar}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
KONTROL = PatternFill("solid", fgColor="FFF4CE")
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def davet_metni(a: Aday, poz: Pozisyon, sirket: str) -> str:
    yer = "[ONLINE GÖRÜŞME BAĞLANTISI]" if katla(poz.format).startswith("online") else "[ADRES VE KAT BİLGİSİ]"
    return (f"Konu: {poz.ad} pozisyonu mülakat daveti\n\nSayın {a.ad},\n\n{sirket} bünyesindeki {poz.ad} pozisyonu başvurunuz için sizi mülakata "
            f"davet etmekten memnuniyet duyarız.\n\nTarih: {a.baslangic:%d.%m.%Y} {GUNLER[a.baslangic.weekday()]}\nSaat: {a.baslangic:%H:%M} – "
            f"{a.bitis:%H:%M}\nFormat: {poz.format}\nYer / bağlantı: {yer}\n\nKatılım durumunuzu bu e-postayı yanıtlayarak bildirmenizi rica ederiz. "
            f"Tarih uygun değilse alternatif zamanlarınızı iletebilirsiniz.\n\nSaygılarımızla,\n[AD SOYAD]\nİnsan Kaynakları")


def ics_yaz(yol: Path, yerlesen: list[Aday], pozisyonlar: dict) -> None:
    s = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//Workers Workless//Mulakat Takvimi//TR", "CALSCALE:GREGORIAN", "X-WR-TIMEZONE:Europe/Istanbul"]
    damga = datetime.now().strftime("%Y%m%dT%H%M%SZ")
    for a in yerlesen:
        poz = pozisyonlar[katla(a.pozisyon)]
        s += ["BEGIN:VEVENT", f"UID:{a.no}-{a.baslangic:%Y%m%d%H%M}@workers-workless", f"DTSTAMP:{damga}",
              f"DTSTART;TZID=Europe/Istanbul:{a.baslangic:%Y%m%dT%H%M%S}", f"DTEND;TZID=Europe/Istanbul:{a.bitis:%Y%m%dT%H%M%S}",
              f"SUMMARY:Mülakat: {a.ad} ({poz.ad})", f"DESCRIPTION:Panel: {', '.join(m.ad for m in a.panel)} · Format: {poz.format}", "END:VEVENT"]
    s.append("END:VCALENDAR")
    yol.write_text("\r\n".join(s) + "\r\n", encoding="utf-8")


def rapor_yaz(cikti: Path, s: dict, mulakatcilar: list[Mulakatci], pozisyonlar: dict, sirket: str) -> Path:
    wb = Workbook()
    tk = wb.active
    tk.title = "Takvim"
    _baslik(tk, ["Tarih", "Gün", "Başlangıç", "Bitiş", "Aday No", "Aday", "Pozisyon", "Format", "Panel", "E-posta", "Davet Gönderildi", "Aday Onayı"],
            (11, 11, 9, 9, 8, 18, 20, 10, 40, 24, 12, 12))
    for a in s["yerlesen"]:
        poz = pozisyonlar[katla(a.pozisyon)]
        tk.append([a.baslangic.date(), GUNLER[a.baslangic.weekday()], a.baslangic.strftime("%H:%M"), a.bitis.strftime("%H:%M"), a.no, a.ad, poz.ad,
                   poz.format, ", ".join(f"{m.ad} ({m.rol})" for m in a.panel), a.eposta, "", ""])
        tk.cell(tk.max_row, 1).number_format = "DD.MM.YYYY"
        tk.cell(tk.max_row, 11).fill = tk.cell(tk.max_row, 12).fill = KONTROL

    mp = wb.create_sheet("Mülakatçı Programı")
    _baslik(mp, ["Mülakatçı", "Rol", "Tarih", "Başlangıç", "Bitiş", "Aday", "Pozisyon"], (20, 10, 11, 9, 9, 18, 20))
    for m in mulakatcilar:
        for bas, bit, a in m.dolu:
            mp.append([m.ad, m.rol, bas.date(), bas.strftime("%H:%M"), bit.strftime("%H:%M"), a.ad, a.pozisyon])
            mp.cell(mp.max_row, 3).number_format = "DD.MM.YYYY"
    mp.append([])
    mp.append(["Yük özeti"])
    mp.cell(mp.max_row, 1).font = Font(bold=True)
    for m in mulakatcilar:
        musait = sum((b - a).total_seconds() / 60 for a, b in m.musait)
        dolu = sum((b - a).total_seconds() / 60 for a, b, _ in m.dolu)
        mp.append([m.ad, m.rol, f"{len(m.dolu)} mülakat", f"{dolu:.0f} dk", f"müsait {musait:.0f} dk", f"doluluk %{dolu / musait * 100:.0f}" if musait else ""])

    ye = wb.create_sheet("Yerleşemeyenler")
    _baslik(ye, ["Aday No", "Aday", "Pozisyon", "Adayın Müsaitliği", "Neden", "Aksiyon"], (8, 18, 20, 40, 70, 26))
    for a in s["yerlesemeyen"]:
        ye.append([a.no, a.ad, a.pozisyon, "; ".join(f"{x:%d.%m.%Y %H:%M}-{y:%H:%M}" for x, y in a.musait) or "—", a.neden, ""])
        ye.cell(ye.max_row, 5).alignment = UST
        ye.cell(ye.max_row, 6).fill = KONTROL

    dv = wb.create_sheet("Davet Metinleri")
    _baslik(dv, ["Aday No", "E-posta", "Davet Metni (taslak)"], (8, 26, 100))
    for a in s["yerlesen"]:
        dv.append([a.no, a.eposta, davet_metni(a, pozisyonlar[katla(a.pozisyon)], sirket)])
        dv.cell(dv.max_row, 3).alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    ics = cikti.with_suffix(".ics")
    ics_yaz(ics, s["yerlesen"], pozisyonlar)
    return ics


def calistir(aday_yolu: Path, mulakatci_yolu: Path, pozisyon_yolu: Path, cikti: Path, tampon_dk: int = 15, sirket: str = "[ŞİRKET ADI]") -> dict:
    pozisyonlar = oku_pozisyonlar(pozisyon_yolu)
    mulakatcilar = oku_mulakatcilar(mulakatci_yolu)
    adaylar = oku_adaylar(aday_yolu)
    s = planla(adaylar, mulakatcilar, pozisyonlar, tampon_dk)
    ics = rapor_yaz(cikti, s, mulakatcilar, pozisyonlar, sirket)
    return {**s, "mulakatcilar": mulakatcilar, "pozisyonlar": pozisyonlar, "ics": ics}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Aday ve mülakatçı müsaitliklerini eşleştirip çakışmasız mülakat takvimi ve davet listesi çıkarır.")
    p.add_argument("--adaylar", type=Path, default=ORNEK / "adaylar.csv", help="Aday No, Ad Soyad, Pozisyon, E-posta, Müsaitlik")
    p.add_argument("--mulakatcilar", type=Path, default=ORNEK / "mulakatcilar.csv", help="Mülakatçı, Rol, Pozisyonlar, Günlük En Çok, Müsaitlik")
    p.add_argument("--pozisyonlar", type=Path, default=ORNEK / "pozisyonlar.csv", help="Pozisyon, Gerekli Roller, Süre (dk), Format")
    p.add_argument("--tampon", type=int, default=15, help="Bir mülakatçının iki mülakatı arasındaki en az süre (dk, varsayılan 15)")
    p.add_argument("--sirket", default="[ŞİRKET ADI]", help="Davet metinlerinde kullanılacak şirket adı")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "mulakat_takvimi.xlsx")
    a = p.parse_args(argv)
    for y in (a.adaylar, a.mulakatcilar, a.pozisyonlar):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        s = calistir(a.adaylar, a.mulakatcilar, a.pozisyonlar, a.cikti, a.tampon, a.sirket)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    for x in s["yerlesen"]:
        print(f"[OK] {x.baslangic:%d.%m.%Y %H:%M}-{x.bitis:%H:%M} {x.no} {x.ad} ({x.pozisyon}) · {', '.join(m.ad for m in x.panel)}")
    for x in s["yerlesemeyen"]:
        print(f"[!] {x.no} {x.ad}: {x.neden}")
    print(f"[OK] Rapor: {a.cikti.resolve()} · takvim: {s['ics'].resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
