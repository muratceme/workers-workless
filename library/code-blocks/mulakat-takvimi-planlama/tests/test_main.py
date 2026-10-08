import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "m.xlsx"
        cls.s = main.calistir(ORNEK / "adaylar.csv", ORNEK / "mulakatcilar.csv", ORNEK / "pozisyonlar.csv", cls.cikti)
        cls.y = {a.no: a for a in cls.s["yerlesen"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_takvim(self):
        self.assertEqual({k: (a.baslangic, [m.ad for m in a.panel]) for k, a in self.y.items()},
                         {"A-08": (datetime(2026, 10, 13, 9, 0), ["İK Uzmanı A"]),
                          "A-04": (datetime(2026, 10, 13, 9, 45), ["İK Uzmanı A", "Yazılım Ekip Lideri"]),
                          "A-01": (datetime(2026, 10, 13, 14, 0), ["İK Uzmanı A", "Finans Müdürü"]),
                          "A-02": (datetime(2026, 10, 13, 15, 0), ["İK Uzmanı A", "Finans Müdürü"]),
                          "A-06": (datetime(2026, 10, 14, 10, 0), ["İK Uzmanı A", "Kıdemli Geliştirici"]),
                          "A-07": (datetime(2026, 10, 14, 13, 0), ["İK Uzmanı B"]),
                          "A-05": (datetime(2026, 10, 15, 14, 0), ["İK Uzmanı A", "Yazılım Ekip Lideri"])})

    def test_cakisma_yok(self):
        for m in self.s["mulakatcilar"]:
            for (a1, b1, _), (a2, b2, _) in zip(m.dolu, m.dolu[1:]):
                self.assertGreaterEqual((a2 - b1).total_seconds(), 15 * 60, m.ad)

    def test_yerlesemeyen(self):
        n = {a.no: a.neden for a in self.s["yerlesemeyen"]}
        self.assertIn("Yönetici rolünde", n["A-03"])
        self.assertIn("İK rolünde", n["A-09"])

    def test_ciktilar(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Takvim", "Mülakatçı Programı", "Yerleşemeyenler", "Davet Metinleri"])
        self.assertIn("Tarih: 13.10.2026 Salı", wb["Davet Metinleri"]["C2"].value)
        ics = self.s["ics"].read_text(encoding="utf-8")
        self.assertEqual(ics.count("BEGIN:VEVENT"), 7)
        self.assertIn("DTSTART;TZID=Europe/Istanbul:20261013T090000", ics)


class KuralTesti(unittest.TestCase):
    def test_araliklar(self):
        self.assertEqual(main.araliklar("13.10.2026 09:00-10:30; 2026-10-14 13.00 – 14.00; bozuk"),
                         [(datetime(2026, 10, 13, 9), datetime(2026, 10, 13, 10, 30)), (datetime(2026, 10, 14, 13), datetime(2026, 10, 14, 14))])

    def test_gunluk_sinir_ve_ortak_bosluk(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t / "p.csv").write_text("Pozisyon;Gerekli Roller;Süre (dk)\nX;İK, Teknik;60\nY;İK;30\n", encoding="utf-8")
            (t / "m.csv").write_text("Mülakatçı;Rol;Pozisyonlar;Günlük En Çok;Müsaitlik\nI;İK;Tümü;1;13.10.2026 09:00-17:00\n"
                                     "T;Teknik;X;;13.10.2026 13:00-14:00\n", encoding="utf-8")
            (t / "a.csv").write_text("Aday No;Ad Soyad;Pozisyon;Müsaitlik\n1;A;Y;13.10.2026 09:00-09:30\n2;B;Y;13.10.2026 10:00-12:00\n"
                                     "3;C;X;13.10.2026 09:00-12:00\n4;D;Z;13.10.2026 09:00-12:00\n", encoding="utf-8")
            s = main.calistir(t / "a.csv", t / "m.csv", t / "p.csv", t / "o.xlsx")
            n = {a.no: a.neden for a in s["yerlesemeyen"]}
            self.assertEqual([a.no for a in s["yerlesen"]], ["1"])
            self.assertIn("İK rolünde", n["2"])                       # günlük sınır 1 doldu
            self.assertIn("Teknik", n["3"])
            self.assertIn("pozisyon tablosunda yok", n["4"])

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--adaylar", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
