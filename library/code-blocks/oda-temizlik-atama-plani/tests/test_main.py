import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def o(no, tur="K", giris=False, ek=0):
    return {"oda": str(no), "kat": str(no)[0], "tur": "Çıkış" if tur == "Ç" else "Konaklama", "giris": giris, "ek": ek, "tip": ""}


def g(*kapasite):
    return [{"ad": f"G{i}", "kapasite": k} for i, k in enumerate(kapasite, 1)]


ODALAR = lambda: [o(204, "Ç", ek=15), o(101, "Ç"), o(102), o(103), o(104, "Ç", True), o(201, "Ç"), o(202)]   # noqa: E731  toplam 195 dk


def dagilim(r):
    d = {}
    for a in r["atamalar"]:
        d.setdefault(a["gorevli"], []).append(a["oda"])
    return d


class AtamaTesti(unittest.TestCase):
    def test_iki_gorevli(self):
        r = main.ata(ODALAR(), g(420, 420))
        self.assertEqual(dagilim(r), {"G1": ["104", "101", "102", "103"], "G2": ["201", "204", "202"]})
        self.assertEqual([(x["gorevli"], x["oda"], x["cikis"], x["konaklama"], x["oncelikli"], x["sure"], x["yuk"], x["katlar"]) for x in r["ozet"]],
                         [("G1", 4, 2, 2, 1, 100, D("23.81"), ["1"]), ("G2", 3, 2, 1, 0, 95, D("22.62"), ["2"])])
        self.assertEqual([(a["sira"], a["oda"], a["tur"], a["oncelik"], a["sure"], a["kumulatif"]) for a in r["atamalar"]][:4],
                         [(1, "104", "Çıkış", "Giriş var", 30, 30), (2, "101", "Çıkış", "", 30, 60), (3, "102", "Konaklama", "", 20, 80), (4, "103", "Konaklama", "", 20, 100)])
        self.assertEqual((r["genel"], r["uyarilar"]), ({"oda": 7, "cikis": 4, "konaklama": 3, "oncelikli": 1, "sure": 195, "kapasite": 840, "yuk": D("23.21"), "gorevli": 2}, []))

    def test_kapasiteyle_orantili_ve_sureler(self):
        r = main.ata(ODALAR(), g(200, 100, 100))
        self.assertEqual((dagilim(r), [x["sure"] for x in r["ozet"]]), ({"G1": ["104", "101", "102", "103"], "G2": ["201", "202"], "G3": ["204"]}, [100, 50, 45]))
        r = main.ata(ODALAR(), g(420, 420), cikis_dk=40, konaklama_dk=10)                  # 101:40 102:10 103:10 104:40 | 201:40 202:10 204:55 → toplam 205
        self.assertEqual((dagilim(r), [x["sure"] for x in r["ozet"]]), ({"G1": ["104", "101", "102", "103"], "G2": ["201", "204", "202"]}, [100, 105]))
        r = main.ata([o(n) for n in (101, 102, 103, 104, 105, 106)], g(100, 100, 100))
        self.assertEqual(dagilim(r), {"G1": ["101", "102"], "G2": ["103", "104"], "G3": ["105", "106"]})
        r = main.ata([o(n, "Ç") | {"kat": ""} for n in (10, 9, 100)], g(1))
        self.assertEqual(dagilim(r), {"G1": ["9", "10", "100"]})                                 # doğal sıralama: 9 < 10 < 100

    def test_uyarilar_ve_hatalar(self):
        r = main.ata(ODALAR(), g(60, 60))
        self.assertEqual((sorted((u["tur"], u["kim"]) for u in r["uyarilar"]), [x["yuk"] for x in r["ozet"]], r["genel"]["yuk"]),
                         ([("Kapasite aşımı", "G1"), ("Kapasite aşımı", "G2"), ("Toplam iş yükü kapasiteyi aşıyor", "Tüm ekip")], [D("166.67"), D("158.33")], D("162.50")))
        r = main.ata([o(101)], g(100, 100))
        self.assertEqual((dagilim(r), [(u["tur"], u["kim"]) for u in r["uyarilar"]]), ({"G1": ["101"]}, [("Oda atanmadı", "G2")]))
        for a, b in (([], g(1)), (ODALAR(), [])):
            with self.assertRaises(ValueError):
                main.ata(a, b)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "odalar.csv", ORNEK / "gorevliler.csv", cikti)
            self.assertEqual(dagilim(r), {"Örnek Görevli A": ["101", "104", "106", "102", "103"], "Örnek Görevli B": ["203", "204", "301", "201", "202", "206"], "Örnek Görevli C": ["302", "303", "304"]})
            self.assertEqual([(x["oda"], x["cikis"], x["oncelikli"], x["sure"], x["yuk"], x["katlar"]) for x in r["ozet"]],
                             [(5, 3, 1, 145, D("82.86"), ["1"]), (6, 3, 1, 150, D("85.71"), ["2", "3"]), (3, 1, 1, 80, D("80.00"), ["3"])])
            self.assertEqual((r["genel"]["oda"], r["genel"]["sure"], r["genel"]["yuk"], r["uyarilar"]), (14, 375, D("83.33"), []))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Görevli Özeti", "Atama Planı", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            oy, gy = Path(tmp) / "o.csv", Path(tmp) / "g.csv"
            oy.write_text("Oda;Kat;Oda Durumu;Giriş;Ek Süre\n101;1;ÇIKIŞ;x;5\n102;1;Stayover;;\n103;1;Temiz;;\n104;1;Kirli;;\n101;1;Konaklama;;\n105;1;Çıkış;;-5\n;;;;\n", encoding="utf-8")
            gy.write_text("Personel;Kapasite\nA;\nB;300\nC;0\n", encoding="utf-8")
            odalar, gorevliler, uy = main.oku(oy, gy, 400)
            self.assertEqual((odalar, gorevliler, sorted(u["tur"] for u in uy)),
                             ([{"oda": "101", "kat": "1", "tur": "Çıkış", "giris": True, "ek": 5, "tip": ""}, {"oda": "102", "kat": "1", "tur": "Konaklama", "giris": False, "ek": 0, "tip": ""}],
                              [{"ad": "A", "kapasite": 400}, {"ad": "B", "kapasite": 300}], ["Görevli satırı okunamadı", "Mükerrer oda", "Oda satırı okunamadı", "Oda satırı okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--odalar", str(oy), "--gorevliler", str(gy), "--cikti", c]
            self.assertEqual(main.main(arg + ["--cikis-dk", "45", "--konaklama-dk", "25", "--kapasite", "300"]), 0)
            for ek in (["--cikis-dk", "0"], ["--kapasite", "-1"]):
                self.assertEqual(main.main(arg + ek), 2)
            self.assertEqual(main.main(["--odalar", str(oy), "--cikti", c]), 2)
            self.assertEqual(main.main(["--odalar", str(oy), "--gorevliler", str(Path(tmp) / "yok.csv")]), 1)
            gy.write_text("Personel\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            oy.write_text("Oda;Kat\n1;1\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
