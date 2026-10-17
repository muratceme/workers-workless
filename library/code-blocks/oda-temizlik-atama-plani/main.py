"""
Oda Temizlik Atama Planı — Gerisi Bende kod bloğu
Kat Hizmetleri › Kat Hizmetleri Müdürü (Housekeeper)

Çıkış ve konaklama odalarını temizlik sürelerine göre kat görevlilerine dengeli ve kat bütünlüğünü koruyarak atar.
  Oda süresi (dk)   = çıkış odası --cikis-dk (30) · konaklama odası --konaklama-dk (20) + odaya özel ek süre     (boş ve arızalı odalar atanmaz)
  Görevli payı      = görevlinin kapasitesi ÷ toplam kapasite                                                   (kapasite verilmezse --kapasite, 420 dk)
  Atama             = odalar kat ve oda numarasına göre sıralanır; sıra bozulmadan görevlilere paylarıyla orantılı dilimler hâlinde dağıtılır
                      (bir odanın orta noktası görevlinin kümülatif pay sınırını aşınca sonraki görevliye geçilir)
  Temizlik sırası   = önce aynı gün girişi olan çıkış odaları, sonra diğer çıkışlar, sonra konaklamalar; her grupta kat ve oda sırasıyla
  Yük %             = görevlinin toplam süresi ÷ kapasitesi × 100;  %100 üstü "kapasite aşımı"
Süreler varsayımdır: oda tipine, standardınıza ve tesise göre değişir; kendi ölçümlerinizle değiştirin. Kod en kısa yürüme yolunu hesaplamaz, kat sırasını korur.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


ODA = {"oda": ("oda no", "oda", "oda numarasi"), "kat": ("kat",), "durum": ("durum", "oda durumu", "temizlik tipi"), "giris": ("giris var", "bugun giris", "giris", "arrival"),
       "ek": ("ek sure", "ek sure dk", "ek dakika"), "tip": ("oda tipi", "tip")}
GOREVLI = {"ad": ("gorevli", "kat gorevlisi", "personel", "ad soyad"), "kapasite": ("kapasite", "kapasite dk", "vardiya dk", "calisma suresi", "calisma suresi dk")}
CIKIS = {"cikis", "c o", "co", "check out", "departure", "ayrilis"}
KONAKLAMA = {"konaklama", "stay", "stayover", "dolu", "konaklayan"}
ATLA = {"bos", "temiz", "vacant", "bos temiz", "arizali", "ooo", "out of order", "kapali", "satis disi"}
EVET = {"evet", "e", "var", "x", "1", "true", "yes"}


def dogal(x: str) -> tuple:
    return tuple((0, int(p)) if p.isdigit() else (1, p) for p in re.findall(r"\d+|\D+", katla(x)))


def oku(oda_yol: Path, gorevli_yol: Path, kapasite: int = 420) -> tuple[list[dict], list[dict], list[dict]]:
    uy, odalar, gorevliler, gorulen = [], [], [], set()
    for r in kayitlar(oda_yol, ODA, ("oda", "durum")):
        no = metin(r["oda"])
        if isinstance(r["oda"], float) and r["oda"] == int(r["oda"]):
            no = str(int(r["oda"]))
        if not no:
            continue
        d, ek = katla(r["durum"]), para(r["ek"]) if metin(r["ek"]) else Decimal(0)
        if d in ATLA:
            continue
        tur = "Çıkış" if d in CIKIS else "Konaklama" if d in KONAKLAMA else None
        if tur is None or ek is None or ek < 0:
            uy.append(uyari("Yüksek", "Oda satırı okunamadı", f"{no} (satır {r['_satir']})", f"Durum '{metin(r['durum'])}', ek süre '{metin(r['ek'])}'. Durum Çıkış, Konaklama, Boş veya Arızalı olmalı. Oda atanmadı."))
            continue
        if katla(no) in gorulen:
            uy.append(uyari("Orta", "Mükerrer oda", no, f"Oda listede birden fazla kez var (satır {r['_satir']}). İlk satır kullanıldı."))
            continue
        gorulen.add(katla(no))
        kat = metin(r["kat"])
        if isinstance(r["kat"], float) and r["kat"] == int(r["kat"]):
            kat = str(int(r["kat"]))
        odalar.append({"oda": no, "kat": kat, "tur": tur, "giris": katla(r["giris"]) in EVET, "ek": int(ek), "tip": metin(r["tip"])})
    for r in kayitlar(gorevli_yol, GOREVLI, ("ad",)):
        ad = metin(r["ad"])
        if not ad:
            continue
        k = para(r["kapasite"]) if metin(r["kapasite"]) else Decimal(kapasite)
        if k is None or k <= 0:
            uy.append(uyari("Yüksek", "Görevli satırı okunamadı", ad, f"Kapasite '{metin(r['kapasite'])}' pozitif dakika olmalı. Görevliye oda atanmadı."))
            continue
        gorevliler.append({"ad": ad, "kapasite": int(k)})
    return odalar, gorevliler, uy


def ata(odalar: list[dict], gorevliler: list[dict], cikis_dk: int = 30, konaklama_dk: int = 20) -> dict:
    if not odalar:
        raise ValueError("Temizlenecek oda yok (çıkış veya konaklama durumunda oda bulunamadı).")
    if not gorevliler:
        raise ValueError("Atama yapılacak kat görevlisi yok.")
    uy = []
    sirali = sorted(({**o, "sure": (cikis_dk if o["tur"] == "Çıkış" else konaklama_dk) + o["ek"]} for o in odalar), key=lambda o: (dogal(o["kat"]), dogal(o["oda"])))
    toplam, top_kap = sum(o["sure"] for o in sirali), sum(g["kapasite"] for g in gorevliler)
    sinir, kum = [], 0
    for g in gorevliler:
        kum += g["kapasite"]
        sinir.append(Decimal(toplam) * kum / top_kap)
    plan, i, gecen = [[] for _ in gorevliler], 0, 0
    for o in sirali:
        while i < len(gorevliler) - 1 and gecen + Decimal(o["sure"]) / 2 > sinir[i]:
            i += 1
        plan[i].append(o)
        gecen += o["sure"]
    atamalar, ozet = [], []
    for g, liste in zip(gorevliler, plan):
        liste.sort(key=lambda o: (0 if o["tur"] == "Çıkış" and o["giris"] else 1 if o["tur"] == "Çıkış" else 2, dogal(o["kat"]), dogal(o["oda"])))
        sure, kum = sum(o["sure"] for o in liste), 0
        for n, o in enumerate(liste, 1):
            kum += o["sure"]
            atamalar.append({"gorevli": g["ad"], "sira": n, "oda": o["oda"], "kat": o["kat"], "tip": o["tip"], "tur": o["tur"], "oncelik": "Giriş var" if o["tur"] == "Çıkış" and o["giris"] else "", "sure": o["sure"], "kumulatif": kum})
        katlar = sorted({o["kat"] for o in liste if o["kat"]}, key=dogal)
        yuk = yuvarla(Decimal(sure) / g["kapasite"] * 100)
        ozet.append({"gorevli": g["ad"], "oda": len(liste), "cikis": sum(1 for o in liste if o["tur"] == "Çıkış"), "konaklama": sum(1 for o in liste if o["tur"] == "Konaklama"),
                     "oncelikli": sum(1 for o in liste if o["tur"] == "Çıkış" and o["giris"]), "sure": sure, "kapasite": g["kapasite"], "yuk": yuk, "katlar": katlar})
        if sure > g["kapasite"]:
            uy.append(uyari("Orta", "Kapasite aşımı", g["ad"], f"Atanan süre {sure} dk, kapasite {g['kapasite']} dk (%{yuk}). Oda devredin veya süreleri gözden geçirin."))
        if not liste:
            uy.append(uyari("Bilgi", "Oda atanmadı", g["ad"], "İş yükü diğer görevlilere sığdığı için bu görevliye oda düşmedi."))
    if toplam > top_kap:
        uy.append(uyari("Yüksek", "Toplam iş yükü kapasiteyi aşıyor", "Tüm ekip", f"Toplam temizlik süresi {toplam} dk, ekip kapasitesi {top_kap} dk. {toplam - top_kap} dk açık var; ek görevli, fazla mesai veya konaklama odalarında hafif servis gerekir."))
    genel = {"oda": len(sirali), "cikis": sum(1 for o in sirali if o["tur"] == "Çıkış"), "konaklama": sum(1 for o in sirali if o["tur"] == "Konaklama"), "oncelikli": sum(1 for o in sirali if o["tur"] == "Çıkış" and o["giris"]),
             "sure": toplam, "kapasite": top_kap, "yuk": yuvarla(Decimal(toplam) / top_kap * 100), "gorevli": len(gorevliler)}
    return {"atamalar": atamalar, "ozet": ozet, "genel": genel, "ayar": (cikis_dk, konaklama_dk), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb, g = Workbook(), r["genel"]
    ws = wb.active
    ws.title = "Görevli Özeti"
    _baslik(ws, ["Kat Görevlisi", "Oda Sayısı", "Çıkış", "Konaklama", "Öncelikli (giriş var)", "Toplam Süre (dk)", "Kapasite (dk)", "Yük %", "Katlar"], (28, 9, 8, 10, 12, 12, 11, 9, 24))
    for o in r["ozet"]:
        _satir(ws, [o["gorevli"], o["oda"], o["cikis"], o["konaklama"], o["oncelikli"] or None, o["sure"], o["kapasite"], o["yuk"], ", ".join(o["katlar"])], para_s=(8,), dolgu={8: KIRMIZI if o["yuk"] > 100 else YESIL})
    n = _satir(ws, ["Toplam", g["oda"], g["cikis"], g["konaklama"], g["oncelikli"] or None, g["sure"], g["kapasite"], g["yuk"], ""], para_s=(8,))
    for c in ws[n]:
        c.font = Font(bold=True)
    ws.append([])
    ws.append([f"Süre varsayımı: çıkış odası {r['ayar'][0]} dk, konaklama odası {r['ayar'][1]} dk, artı odaya özel ek süre. Atama kat sırasını korur; son dakika değişikliklerini siz düzeltirsiniz."])
    ws = wb.create_sheet("Atama Planı")
    _baslik(ws, ["Kat Görevlisi", "Sıra", "Oda No", "Kat", "Oda Tipi", "Temizlik Tipi", "Öncelik", "Süre (dk)", "Kümülatif (dk)", "Tamamlandı (siz doldurun)"], (28, 6, 10, 8, 16, 13, 12, 10, 12, 20))
    for a in r["atamalar"]:
        _satir(ws, [a["gorevli"], a["sira"], a["oda"], a["kat"], a["tip"], a["tur"], a["oncelik"], a["sure"], a["kumulatif"], ""], dolgu={7: SARI if a["oncelik"] else None, 6: MAVI if a["tur"] == "Çıkış" else None})
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(oda_yol: Path, gorevli_yol: Path, cikti: Path, kapasite: int = 420, **ayar) -> dict:
    odalar, gorevliler, uy = oku(oda_yol, gorevli_yol, kapasite)
    r = ata(odalar, gorevliler, **ayar)
    uy += r["uyarilar"]
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Çıkış ve konaklama odalarını kat görevlilerine dengeli biçimde atar.")
    p.add_argument("--odalar", type=Path, default=ORNEK / "odalar.csv", help="Oda No, Kat, Durum (Çıkış / Konaklama / Boş / Arızalı), Giriş Var, Ek Süre (dk), Oda Tipi")
    p.add_argument("--gorevliler", type=Path, help="Görevli, Kapasite (dk)")
    p.add_argument("--cikis-dk", type=int, default=30, help="Çıkış odası temizlik süresi, dk (varsayılan 30)")
    p.add_argument("--konaklama-dk", type=int, default=20, help="Konaklama odası temizlik süresi, dk (varsayılan 20)")
    p.add_argument("--kapasite", type=int, default=420, help="Görevli dosyasında kapasite yoksa kullanılacak dakika (varsayılan 420)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "temizlik_atama_plani.xlsx")
    a = p.parse_args(argv)
    gorevli = a.gorevliler or (ORNEK / "gorevliler.csv" if a.odalar == ORNEK / "odalar.csv" else None)
    if gorevli is None or a.cikis_dk <= 0 or a.konaklama_dk <= 0 or a.kapasite <= 0:
        print("[X] --gorevliler dosyası verilmeli; --cikis-dk, --konaklama-dk ve --kapasite pozitif tam sayı olmalı")
        return 2
    for y in (a.odalar, gorevli):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.odalar, gorevli, a.cikti, a.kapasite, cikis_dk=a.cikis_dk, konaklama_dk=a.konaklama_dk)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    g = r["genel"]
    print(f"[OK] {g['oda']} oda ({g['cikis']} çıkış, {g['konaklama']} konaklama) · {g['gorevli']} görevli · toplam {g['sure']} dk, kapasite {g['kapasite']} dk (%{g['yuk']})")
    for o in r["ozet"]:
        print(f"[{'!' if o['yuk'] > 100 else 'OK'}] {o['gorevli']}: {o['oda']} oda, {o['sure']} dk (%{o['yuk']}), kat {', '.join(o['katlar']) or '-'}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
