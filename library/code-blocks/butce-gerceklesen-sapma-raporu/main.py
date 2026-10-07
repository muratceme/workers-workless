"""
Bütçe-Gerçekleşen Sapma Raporu — Workers / Workless kod bloğu
Finans › Bütçe ve Raporlama Uzmanı

Yıllık bütçeyi (hesap × masraf merkezi × ay) muhasebeden alınan gerçekleşen kayıtlarla karşılaştırır.
Rapor ayı ve yılbaşından bugüne (YTD) sapmayı TL ve % olarak, lehte/aleyhte yönüyle verir; önemlilik
eşiğini (% ve TL birlikte) aşan sapmaları açıklama alanıyla işaretler, yıl sonu tahminini çıkarır.

- Gerçekleşen, ayrıntılı hesaplardan (770.02.001) bütçe satırına (770.02) en uzun önek eşleşmesiyle toplanır.
- Gelir/gider yönü Tekdüzen Hesap Planı'ndan belirlenir (60, 64, 67 gelir; 61-63, 65, 66, 68, 69 ve 7'li
  maliyet hesapları gider); "Tür" sütunu verilirse o kullanılır.
- Bütçede olmayan hesaplara yapılan kayıtlar "Bütçelenmemiş" olarak ayrıca listelenir.
İnternete bağlanmaz.

Kullanım:
    python main.py                                         # örnek bütçe ve muavinle dener (Eylül 2026)
    python main.py --butce butce_2026.xlsx --gerceklesen muavin.xlsx --ay 9
    python main.py --butce butce.xlsx --gerceklesen muavin.xlsx --esik-yuzde 10 --esik-tutar 25000
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
AYLAR = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]
GELIR_GRUPLARI = ("60", "64", "67")        # Tekdüzen: Brüt satışlar, diğer faaliyet gelirleri, olağandışı gelirler


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def para(x) -> Decimal:
    if x in (None, ""):
        return Decimal(0)
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("₺", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return Decimal(0)


def tarih(x) -> date | None:
    if x in (None, ""):
        return None
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(str(x).strip(), f).date()
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ayirici = csv.Sniffer().sniff(metin.splitlines()[0], delimiters=",;\t").delimiter
        satirlar = list(csv.reader(metin.splitlines(), delimiter=ayirici))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


ALANLAR = {
    "hesap": ("hesap kodu", "hesap", "hesap no", "kod"),
    "hesap_adi": ("hesap adı", "hesap adi", "açıklama", "kalem", "bütçe kalemi"),
    "mm": ("masraf merkezi", "maliyet merkezi", "departman", "mm"),
    "tur": ("tür", "tip", "gelir/gider"),
    "ay": ("ay", "dönem", "donem"),
    "tarih": ("tarih", "fiş tarihi", "belge tarihi", "işlem tarihi"),
    "tutar": ("tutar", "bütçe", "gerçekleşen", "bütçe tutarı"),
    "borc": ("borç", "borc"),
    "alacak": ("alacak",),
}


def sutunlar(baslik: list) -> dict:
    b = [kucuk(x) for x in baslik]
    return {alan: next((i for i, x in enumerate(b) if x in adlar), None) for alan, adlar in ALANLAR.items()}


def ay_sutunlari(baslik: list) -> dict[int, int]:
    """Geniş biçim: Ocak … Aralık (veya 1 … 12 / 01 … 12) başlıklı sütunlar."""
    b = [kucuk(x) for x in baslik]
    adlar = [kucuk(a) for a in AYLAR]
    sonuc = {}
    for i, x in enumerate(b):
        if x in adlar:
            sonuc[adlar.index(x) + 1] = i
        elif re.fullmatch(r"0?([1-9]|1[0-2])", x):
            sonuc[int(x)] = i
    return sonuc if len(sonuc) >= 2 else {}


def ay_coz(x) -> int | None:
    if x in (None, ""):
        return None
    t = tarih(x)
    if t:
        return t.month
    s = kucuk(x)
    for i, a in enumerate(AYLAR, 1):
        if s.startswith(kucuk(a)):
            return i
    m = re.search(r"(?:^|\D)(0?[1-9]|1[0-2])$", s) or re.match(r"^(0?[1-9]|1[0-2])(?:\D|$)", s)
    return int(m.group(1)) if m else None


def gelir_mi(hesap: str, tur: str = "") -> bool:
    t = kucuk(tur)
    if t.startswith("gelir"):
        return True
    if t.startswith("gider") or t.startswith("maliyet"):
        return False
    return re.sub(r"\D", "", hesap)[:2] in GELIR_GRUPLARI


def kod(hesap) -> str:
    return re.sub(r"\s", "", str(hesap or "")).strip(".")


# ----------------------------------------------------------------------------
# Okuma
# ----------------------------------------------------------------------------

@dataclass
class Satir:
    hesap: str
    ad: str
    mm: str
    gelir: bool
    butce: dict = field(default_factory=lambda: defaultdict(Decimal))     # ay → tutar
    fiili: dict = field(default_factory=lambda: defaultdict(Decimal))


def butce_oku(yol: Path) -> tuple[dict, bool]:
    satirlar = tablo_oku(yol)
    k, aylar = sutunlar(satirlar[0]), ay_sutunlari(satirlar[0])
    if k["hesap"] is None:
        raise SystemExit(f"Bütçe dosyasında 'Hesap Kodu' sütunu bulunamadı. Başlıklar: {satirlar[0]}")
    if not aylar and (k["ay"] is None or k["tutar"] is None):
        raise SystemExit("Bütçe ya Ocak … Aralık sütunlarıyla (geniş) ya da Ay + Tutar sütunlarıyla (uzun) olmalı.")
    al = lambda r, a: r[k[a]] if k[a] is not None and k[a] < len(r) else None  # noqa: E731
    butce: dict[tuple, Satir] = {}
    for r in satirlar[1:]:
        h = kod(al(r, "hesap"))
        if not h:
            continue
        mm = str(al(r, "mm") or "").strip()
        s = butce.setdefault((h, mm), Satir(h, str(al(r, "hesap_adi") or "").strip(), mm, gelir_mi(h, str(al(r, "tur") or ""))))
        if aylar:
            for ay, i in aylar.items():
                s.butce[ay] += para(r[i] if i < len(r) else None)
        else:
            ay = ay_coz(al(r, "ay"))
            if ay:
                s.butce[ay] += para(al(r, "tutar"))
    return butce, k["mm"] is not None


def gerceklesen_oku(yol: Path, yil: int | None) -> tuple[list[tuple], bool, set[int]]:
    """(hesap, mm, ay, tutar, ad, tür) listesi. Borç/Alacak verilmişse yön hesap türünden belirlenir."""
    satirlar = tablo_oku(yol)
    k, aylar = sutunlar(satirlar[0]), ay_sutunlari(satirlar[0])
    if k["hesap"] is None:
        raise SystemExit(f"Gerçekleşen dosyasında 'Hesap Kodu' sütunu bulunamadı. Başlıklar: {satirlar[0]}")
    al = lambda r, a: r[k[a]] if k[a] is not None and k[a] < len(r) else None  # noqa: E731
    kayitlar, yillar = [], set()
    for r in satirlar[1:]:
        h = kod(al(r, "hesap"))
        if not h:
            continue
        mm, ad, tur = str(al(r, "mm") or "").strip(), str(al(r, "hesap_adi") or "").strip(), str(al(r, "tur") or "")
        if aylar:
            for ay, i in aylar.items():
                kayitlar.append((h, mm, ay, para(r[i] if i < len(r) else None), ad, tur))
            continue
        t = tarih(al(r, "tarih")) if k["tarih"] is not None else None
        if t:
            yillar.add(t.year)
            if yil and t.year != yil:
                continue
        ay = t.month if t else ay_coz(al(r, "ay"))
        if not ay:
            continue
        if k["borc"] is not None and k["alacak"] is not None:
            net = para(al(r, "borc")) - para(al(r, "alacak"))
            tutar = -net if gelir_mi(h, tur) else net
        else:
            tutar = para(al(r, "tutar"))
        kayitlar.append((h, mm, ay, tutar, ad, tur))
    return kayitlar, k["mm"] is not None, yillar


def onek_bul(hesap: str, adaylar: list[str]) -> str | None:
    """En uzun önek eşleşmesi: 770.02.001 → 770.02 (770.0 değil, hane sınırında)."""
    en = None
    for a in adaylar:
        if hesap == a or (hesap.startswith(a) and not hesap[len(a)].isalnum()):
            if en is None or len(a) > len(en):
                en = a
    return en


# ----------------------------------------------------------------------------
# Hesap
# ----------------------------------------------------------------------------

def sapma_yorumu(gelir: bool, fark: Decimal) -> str:
    if fark == 0:
        return "—"
    lehte = fark > 0 if gelir else fark < 0
    return "Lehte" if lehte else "Aleyhte"


def yuzde(fark: Decimal, butce: Decimal) -> float | None:
    return None if butce == 0 else float(fark / abs(butce) * 100)


def onemli(fark: Decimal, butce: Decimal, esik_yuzde: Decimal, esik_tutar: Decimal) -> bool:
    if abs(fark) < esik_tutar:
        return False
    return butce == 0 or abs(fark) / abs(butce) * 100 >= esik_yuzde


def calistir(butce_yolu: Path, gerceklesen_yolu: Path, cikti: Path, ay: int | None = None, yil: int | None = None,
             esik_yuzde: Decimal = Decimal(10), esik_tutar: Decimal = Decimal(10000)) -> dict:
    butce, butce_mm = butce_oku(butce_yolu)
    kayitlar, fiili_mm, yillar = gerceklesen_oku(gerceklesen_yolu, yil)
    uyarilar = []
    if yil is None and len(yillar) > 1:      # birden çok yıl: en son yıl alınır
        yil = max(yillar)
        kayitlar, _, _ = gerceklesen_oku(gerceklesen_yolu, yil)
        uyarilar.append(f"Gerçekleşen dosyasında birden çok yıl var ({', '.join(map(str, sorted(yillar)))}); "
                        f"yalnız {yil} kayıtları alındı (başka yıl için --yil verin)")
    mm_kullan = butce_mm and fiili_mm
    if butce_mm != fiili_mm:
        uyarilar.append("Masraf merkezi yalnız bir dosyada var; karşılaştırma hesap bazında yapıldı")
        birlesik: dict[tuple, Satir] = {}
        for s in butce.values():
            b = birlesik.setdefault((s.hesap, ""), Satir(s.hesap, s.ad, "", s.gelir))
            for a, v in s.butce.items():
                b.butce[a] += v
        butce = birlesik

    hesaplar = sorted({h for h, _ in butce})
    butcelenmemis: dict[tuple, Satir] = {}
    for h, mm, a, tutar, ad, tur in kayitlar:
        mm = mm if mm_kullan else ""
        hedef = onek_bul(h, hesaplar)
        if hedef is not None and (hedef, mm) in butce:
            butce[(hedef, mm)].fiili[a] += tutar
        else:
            anahtar = (hedef or h, mm)
            s = butcelenmemis.setdefault(anahtar, Satir(hedef or h, ad, mm, gelir_mi(h, tur)))
            s.fiili[a] += tutar

    tum_aylar = {a for s in butce.values() for a in s.fiili} | {a for s in butcelenmemis.values() for a in s.fiili}
    if not tum_aylar:
        raise SystemExit("Gerçekleşen dosyasında bütçe yılına ait kayıt bulunamadı.")
    ay = ay or max(tum_aylar)

    satirlar = []
    for s in list(butce.values()) + list(butcelenmemis.values()):
        b_ay, f_ay = s.butce.get(ay, Decimal(0)), s.fiili.get(ay, Decimal(0))
        b_ytd = sum((s.butce.get(a, Decimal(0)) for a in range(1, ay + 1)), Decimal(0))
        f_ytd = sum((s.fiili.get(a, Decimal(0)) for a in range(1, ay + 1)), Decimal(0))
        b_yil = sum(s.butce.values(), Decimal(0))
        tahmin = f_ytd + sum((s.butce.get(a, Decimal(0)) for a in range(ay + 1, 13)), Decimal(0))
        fark_ay, fark_ytd = f_ay - b_ay, f_ytd - b_ytd
        on_ay = onemli(fark_ay, b_ay, esik_yuzde, esik_tutar)
        on_ytd = onemli(fark_ytd, b_ytd, esik_yuzde, esik_tutar)
        not_ = []
        if (s.hesap, s.mm) in butcelenmemis:
            not_.append("Bütçelenmemiş hesap")
        if on_ay and not on_ytd:
            not_.append("Ay sapması YTD'de önemsiz: zamanlama (dönem kayması) olabilir")
        if on_ytd and b_yil and abs(tahmin - b_yil) / abs(b_yil) * 100 >= esik_yuzde:
            not_.append(f"Yıl sonu tahmini bütçeden %{abs(tahmin - b_yil) / abs(b_yil) * 100:.0f} sapıyor")
        satirlar.append({
            "hesap": s.hesap, "ad": s.ad, "mm": s.mm, "gelir": s.gelir,
            "b_ay": b_ay, "f_ay": f_ay, "fark_ay": fark_ay, "yuzde_ay": yuzde(fark_ay, b_ay), "yon_ay": sapma_yorumu(s.gelir, fark_ay),
            "b_ytd": b_ytd, "f_ytd": f_ytd, "fark_ytd": fark_ytd, "yuzde_ytd": yuzde(fark_ytd, b_ytd),
            "yon_ytd": sapma_yorumu(s.gelir, fark_ytd), "b_yil": b_yil, "tahmin": tahmin,
            "gerceklesme": None if not b_yil else float(f_ytd / b_yil * 100),
            "onemli_ay": on_ay, "onemli_ytd": on_ytd, "not": "; ".join(not_), "satir": s,
        })
    satirlar.sort(key=lambda x: (not x["gelir"], x["hesap"], x["mm"]))
    _rapor(satirlar, ay, cikti, esik_yuzde, esik_tutar, uyarilar)
    return {"satirlar": satirlar, "ay": ay, "yil": yil, "uyarilar": uyarilar,
            "onemli": [x for x in satirlar if x["onemli_ay"] or x["onemli_ytd"]]}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
YESIL = PatternFill("solid", fgColor="E3F5E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
PARA = "#,##0.00"
YUZDE = '0.0"%"'


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = f"D{ws.max_row + 1}"


def _rapor(satirlar, ay, cikti, esik_yuzde, esik_tutar, uyarilar):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sapma Raporu"
    _baslik(ws, ["Hesap", "Hesap Adı", "Masraf Merkezi", "Tür",
                 f"{AYLAR[ay - 1]} Bütçe", f"{AYLAR[ay - 1]} Gerçekleşen", "Ay Sapma", "Ay Sapma %", "Ay Yön",
                 "YTD Bütçe", "YTD Gerçekleşen", "YTD Sapma", "YTD Sapma %", "YTD Yön",
                 "Yıllık Bütçe", "Gerçekleşme %", "Yıl Sonu Tahmini", "Önemli Sapma", "Not", "Açıklama (doldurunuz)"],
            (12, 30, 16, 8, 15, 15, 14, 11, 9, 16, 16, 14, 11, 9, 16, 12, 16, 13, 50, 40))
    toplam = {g: defaultdict(Decimal) for g in (True, False)}
    for x in satirlar:
        ws.append([x["hesap"], x["ad"], x["mm"], "Gelir" if x["gelir"] else "Gider",
                   float(x["b_ay"]), float(x["f_ay"]), float(x["fark_ay"]), x["yuzde_ay"], x["yon_ay"],
                   float(x["b_ytd"]), float(x["f_ytd"]), float(x["fark_ytd"]), x["yuzde_ytd"], x["yon_ytd"],
                   float(x["b_yil"]), x["gerceklesme"], float(x["tahmin"]),
                   " / ".join(t for t, v in (("Ay", x["onemli_ay"]), ("YTD", x["onemli_ytd"])) if v), x["not"], ""])
        r = ws.max_row
        for c in (5, 6, 7, 10, 11, 12, 15, 17):
            ws.cell(r, c).number_format = PARA
        for c in (8, 13, 16):
            ws.cell(r, c).number_format = YUZDE
        for c, yon in ((9, x["yon_ay"]), (14, x["yon_ytd"])):
            ws.cell(r, c).fill = YESIL if yon == "Lehte" else KIRMIZI if yon == "Aleyhte" else PatternFill()
        if x["onemli_ay"] or x["onemli_ytd"]:
            ws.cell(r, 18).fill = KIRMIZI
            ws.cell(r, 20).fill = SARI
        for k in ("b_ay", "f_ay", "b_ytd", "f_ytd", "b_yil", "tahmin"):
            toplam[x["gelir"]][k] += x[k]
    ws.auto_filter.ref = ws.dimensions

    o = wb.create_sheet("Özet", 0)
    o.append([f"Bütçe-Gerçekleşen Sapma Raporu · {AYLAR[ay - 1]} ve yılbaşından bugüne (Ocak–{AYLAR[ay - 1]})"])
    o["A1"].font = Font(bold=True, size=13)
    o.append([f"Önemlilik eşiği: sapma ≥ %{esik_yuzde} ve ≥ {esik_tutar:,.0f} TL".replace(",", ".")])
    o.append([])
    o.append(["", f"{AYLAR[ay - 1]} Bütçe", f"{AYLAR[ay - 1]} Gerçekleşen", "Ay Sapma", "YTD Bütçe", "YTD Gerçekleşen",
              "YTD Sapma", "Yıllık Bütçe", "Yıl Sonu Tahmini"])
    for h in o[4]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    g, d = toplam[True], toplam[False]
    for ad, v in (("Gelirler", g), ("Giderler", d),
                  ("Fark (gelir − gider)", {k: g[k] - d[k] for k in ("b_ay", "f_ay", "b_ytd", "f_ytd", "b_yil", "tahmin")})):
        o.append([ad, float(v["b_ay"]), float(v["f_ay"]), float(v["f_ay"] - v["b_ay"]), float(v["b_ytd"]), float(v["f_ytd"]),
                  float(v["f_ytd"] - v["b_ytd"]), float(v["b_yil"]), float(v["tahmin"])])
        for c in range(2, 10):
            o.cell(o.max_row, c).number_format = PARA
    o.cell(o.max_row, 1).font = Font(bold=True)
    o.append([])
    onemliler = [x for x in satirlar if x["onemli_ay"] or x["onemli_ytd"]]
    o.append([f"Önemli sapmalar ({len(onemliler)})"])
    o.cell(o.max_row, 1).font = Font(bold=True)
    o.append(["Hesap", "Hesap Adı", "Masraf Merkezi", "Ay Sapma", "Ay Sapma %", "YTD Sapma", "YTD Sapma %", "Yön (YTD)", "Not"])
    for h in o[o.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for x in sorted(onemliler, key=lambda x: -abs(x["fark_ytd"])):
        o.append([x["hesap"], x["ad"], x["mm"], float(x["fark_ay"]), x["yuzde_ay"], float(x["fark_ytd"]), x["yuzde_ytd"],
                  x["yon_ytd"], x["not"]])
        for c in (4, 6):
            o.cell(o.max_row, c).number_format = PARA
        for c in (5, 7):
            o.cell(o.max_row, c).number_format = YUZDE
        o.cell(o.max_row, 8).fill = YESIL if x["yon_ytd"] == "Lehte" else KIRMIZI if x["yon_ytd"] == "Aleyhte" else PatternFill()
    for j, w in enumerate((22, 30, 18, 16, 16, 16, 16, 16, 60), 1):
        o.column_dimensions[get_column_letter(j)].width = w
    for u in uyarilar:
        o.append(["Uyarı", u])

    # Masraf merkezi özeti (giderler)
    mmler = defaultdict(lambda: defaultdict(Decimal))
    for x in satirlar:
        if not x["gelir"]:
            for k in ("b_ytd", "f_ytd", "b_yil", "tahmin"):
                mmler[x["mm"] or "(tümü)"][k] += x[k]
    m = wb.create_sheet("Masraf Merkezi")
    _baslik(m, ["Masraf Merkezi", "YTD Bütçe", "YTD Gerçekleşen", "YTD Sapma", "YTD Sapma %", "Yıllık Bütçe", "Yıl Sonu Tahmini"],
            (22, 16, 16, 16, 12, 16, 16))
    m.freeze_panes = "B2"
    for ad, v in sorted(mmler.items()):
        fark = v["f_ytd"] - v["b_ytd"]
        m.append([ad, float(v["b_ytd"]), float(v["f_ytd"]), float(fark), yuzde(fark, v["b_ytd"]), float(v["b_yil"]), float(v["tahmin"])])
        for c in (2, 3, 4, 6, 7):
            m.cell(m.max_row, c).number_format = PARA
        m.cell(m.max_row, 5).number_format = YUZDE
    if len(mmler) > 1:
        gr = BarChart()
        gr.title, gr.height, gr.width = "YTD bütçe ve gerçekleşen (giderler)", 8, 18
        gr.add_data(Reference(m, min_col=2, max_col=3, min_row=1, max_row=1 + len(mmler)), titles_from_data=True)
        gr.set_categories(Reference(m, min_col=1, min_row=2, max_row=1 + len(mmler)))
        m.add_chart(gr, "I2")

    # Aylık seyir
    a = wb.create_sheet("Aylık Seyir")
    _baslik(a, ["Ay", "Gelir Bütçe", "Gelir Gerçekleşen", "Gider Bütçe", "Gider Gerçekleşen"], (12, 16, 18, 16, 18))
    a.freeze_panes = "B2"
    for i in range(1, 13):
        satir = [AYLAR[i - 1]]
        for gelir in (True, False):
            sb = sum((x["satir"].butce.get(i, Decimal(0)) for x in satirlar if x["gelir"] == gelir), Decimal(0))
            sf = sum((x["satir"].fiili.get(i, Decimal(0)) for x in satirlar if x["gelir"] == gelir), Decimal(0))
            satir += [float(sb), float(sf) if i <= ay else None]
        a.append(satir)
        for c in range(2, 6):
            a.cell(a.max_row, c).number_format = PARA
    lc = LineChart()
    lc.title, lc.height, lc.width = "Aylık bütçe ve gerçekleşen", 8, 20
    lc.add_data(Reference(a, min_col=2, max_col=5, min_row=1, max_row=13), titles_from_data=True)
    lc.set_categories(Reference(a, min_col=1, min_row=2, max_row=13))
    a.add_chart(lc, "G2")

    b = wb.create_sheet("Bilgi")
    for s in [["Sapma", "Gerçekleşen − Bütçe. Gelirde pozitif sapma lehte, giderde aleyhtedir"],
              ["Önemli sapma", f"|sapma| ≥ {esik_tutar} TL ve |sapma| ≥ bütçenin %{esik_yuzde}'u (bütçe 0 ise yalnız tutar eşiği)"],
              ["Yıl sonu tahmini", "YTD gerçekleşen + kalan ayların bütçesi"],
              ["Hesap eşleşmesi", "Gerçekleşen hesap kodu, bütçedeki en uzun önek hesaba toplanır (770.02.001 → 770.02)"],
              ["Gelir/gider", "Tür sütunu yoksa Tekdüzen Hesap Planı: 60, 64, 67 gelir; diğerleri gider"],
              ["Borç/Alacak", "Gerçekleşen Borç/Alacak ile verilmişse gider = borç − alacak, gelir = alacak − borç"]]:
        b.append(s)
    b.column_dimensions["A"].width = 18
    b.column_dimensions["B"].width = 110
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description="Bütçe ile gerçekleşeni hesap ve masraf merkezi bazında karşılaştırır.")
    ap.add_argument("--butce", type=Path, default=BURASI / "ornek_veri" / "butce_2026.csv",
                    help="Bütçe (.xlsx/.csv): Hesap Kodu, [Hesap Adı], [Masraf Merkezi], Ocak … Aralık (veya Ay + Tutar)")
    ap.add_argument("--gerceklesen", type=Path, default=BURASI / "ornek_veri" / "gerceklesen_muavin.csv",
                    help="Gerçekleşen (.xlsx/.csv): Hesap Kodu, Tarih (veya Ay), Borç + Alacak (veya Tutar), [Masraf Merkezi]")
    ap.add_argument("--ay", type=int, choices=range(1, 13), metavar="1-12", help="Rapor ayı (varsayılan: son gerçekleşen ay)")
    ap.add_argument("--yil", type=int, help="Gerçekleşen dosyasında yalnız bu yılın kayıtlarını al (varsayılan: en son yıl)")
    ap.add_argument("--esik-yuzde", type=Decimal, default=Decimal(10), help="Önemlilik eşiği, %% (varsayılan 10)")
    ap.add_argument("--esik-tutar", type=Decimal, default=Decimal(10000), help="Önemlilik eşiği, TL (varsayılan 10000)")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "butce_sapma.xlsx")
    a = ap.parse_args(argv)
    s = calistir(a.butce, a.gerceklesen, a.cikti, a.ay, a.yil, a.esik_yuzde, a.esik_tutar)
    tl = lambda x: f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")  # noqa: E731
    print(f"[OK] {len(s['satirlar'])} bütçe satırı · rapor ayı {AYLAR[s['ay'] - 1]}")
    for x in sorted(s["onemli"], key=lambda x: -abs(x["fark_ytd"]))[:10]:
        print(f"[!] {x['hesap']} {x['ad'][:28]:<28} {x['mm'][:14]:<14} YTD {tl(x['fark_ytd']):>14} TL  {x['yon_ytd']}"
              + (f"  · {x['not']}" if x["not"] else ""))
    for u in s["uyarilar"]:
        print(f"[!] {u}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
