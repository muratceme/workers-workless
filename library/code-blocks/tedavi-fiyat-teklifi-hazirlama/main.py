"""
Tedavi Fiyat Teklifi Hazırlama — Gerisi Bende kod bloğu
Uluslararası Hasta Hizmetleri › Uluslararası Hasta Koordinatörü

Tedavi planı ve paket fiyat listesinden konaklama ve transfer dahil hastaya gönderilecek teklif dosyasını hazırlar.
  Kalem tutarı   = adet × fiyat listesindeki birim fiyat        (plandaki kalem, fiyat listesinde kod veya ad ile aranır)
  Ara toplam     = kategori bazında (Tedavi, Konaklama, Transfer, Diğer) kalem tutarlarının toplamı
  İndirim        = toplam × --indirim %                        (--indirim-kategori verilirse yalnız o kategorinin ara toplamına)
  Genel toplam   = toplam − indirim
  Geçerlilik     = teklif tarihi + --gecerlilik-gun (varsayılan 30)
Fiyatlar sizin fiyat listenizden gelir; kod fiyat, kur veya vergi bilmez. Tüm kalemler aynı para biriminde olmalıdır; kur çevirisi yapılmaz.
Teklif metni taslaktır: tıbbi uygunluk, tedavi planı ve fiyat hekim ve kurum onayından geçmeden gönderilmemelidir. Hasta adı gerekmez; referans kodu kullanılır.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


PLAN = {"kalem": ("kalem", "kod", "hizmet", "tedavi", "islem", "kalem kodu"), "adet": ("adet", "miktar", "sayi", "gece"), "not": ("not", "aciklama")}
FIYAT = {"kod": ("kod", "kalem kodu", "hizmet kodu"), "ad": ("kalem", "ad", "hizmet", "hizmet adi", "tedavi"), "kategori": ("kategori", "grup"), "fiyat": ("birim fiyat", "fiyat", "paket fiyat"),
         "para": ("para birimi", "doviz", "birim"), "en": ("ingilizce ad", "ingilizce", "name", "english name", "name en")}
KATEGORI = {"tedavi": ("Tedavi", "Treatment"), "konaklama": ("Konaklama", "Accommodation"), "transfer": ("Transfer", "Transfer"), "diger": ("Diğer", "Other")}
METIN = {"tr": {"baslik": "TEDAVİ FİYAT TEKLİFİ", "ref": "Referans", "tarih": "Teklif tarihi", "gecerli": "Geçerlilik", "ara": "Ara toplam", "toplam": "Toplam", "indirim": "İndirim", "genel": "GENEL TOPLAM",
                "notlar": "Notlar", "n1": "Bu teklif ön bilgilendirme amaçlıdır; kesin tedavi planı ve fiyat hekim muayenesinden sonra belirlenir.",
                "n2": "Teklifte yazmayan hizmetler fiyata dahil değildir."},
         "en": {"baslik": "TREATMENT PRICE QUOTE", "ref": "Reference", "tarih": "Quote date", "gecerli": "Valid until", "ara": "Subtotal", "toplam": "Total", "indirim": "Discount", "genel": "GRAND TOTAL",
                "notlar": "Notes", "n1": "This quote is for preliminary information only; the final treatment plan and price are determined after the doctor's examination.",
                "n2": "Services not listed in this quote are not included in the price."}}
Z = Decimal(0)


def tutar(x: Decimal, dil: str) -> str:
    s = f"{x:,.2f}"
    return s if dil == "en" else s.replace(",", "X").replace(".", ",").replace("X", ".")


def oku(plan_yol: Path, fiyat_yol: Path) -> tuple[list[dict], dict, list[dict]]:
    uy, liste, plan = [], {}, []
    for r in kayitlar(fiyat_yol, FIYAT, ("ad", "fiyat")):
        ad = metin(r["ad"])
        if not ad:
            continue
        f = para(r["fiyat"])
        if f is None or f < 0:
            uy.append(uyari("Yüksek", "Liste fiyatı okunamadı", ad, f"{fiyat_yol.name} satır {r['_satir']}: '{metin(r['fiyat'])}'. Kalem fiyat listesine alınmadı."))
            continue
        k = katla(r["kategori"])
        x = {"kod": metin(r["kod"]), "ad": ad, "kategori": k if k in KATEGORI else "diger", "fiyat": f, "para": metin(r["para"]).upper(), "en": metin(r["en"])}
        for a in {katla(x["kod"]), katla(ad)} - {""}:
            if a in liste:
                uy.append(uyari("Orta", "Fiyat listesinde mükerrer kalem", ad, f"'{a}' birden fazla satırda; ilk satırdaki fiyat kullanıldı."))
            liste.setdefault(a, x)
    for r in kayitlar(plan_yol, PLAN, ("kalem",)):
        k = metin(r["kalem"])
        if not k:
            continue
        adet = para(r["adet"]) if metin(r["adet"]) else Decimal(1)
        if adet is None or adet <= 0:
            uy.append(uyari("Yüksek", "Adet okunamadı", k, f"{plan_yol.name} satır {r['_satir']}: '{metin(r['adet'])}'. Kalem teklife alınmadı."))
            continue
        plan.append({"kalem": k, "adet": adet, "not": metin(r["not"])})
    return plan, liste, uy


def teklif(plan: list[dict], liste: dict, teklif_tarihi: date, indirim: Decimal = Z, indirim_kategori: str | None = None, gecerlilik: int = 30) -> dict:
    kalemler, bulunamayan = [], []
    for p in plan:
        x = liste.get(katla(p["kalem"]))
        if x is None:
            bulunamayan.append(p["kalem"])
            continue
        kalemler.append(x | {"adet": p["adet"], "tutar": yuvarla(p["adet"] * x["fiyat"]), "not": p["not"]})
    if not kalemler:
        raise ValueError("Teklife girecek kalem yok: plandaki kalemler fiyat listesinde bulunamadı.")
    paralar = sorted({k["para"] for k in kalemler})
    if len(paralar) > 1:
        raise ValueError(f"Kalemler farklı para birimlerinde: {', '.join(p or '(boş)' for p in paralar)}. Kur çevirisi yapılmaz; fiyat listesini tek para birimine getirin.")
    ara = {k: sum((x["tutar"] for x in kalemler if x["kategori"] == k), Z) for k in KATEGORI if any(x["kategori"] == k for x in kalemler)}
    toplam = sum(ara.values(), Z)
    if indirim_kategori and indirim_kategori not in ara:
        raise ValueError(f"İndirim uygulanacak kategori teklifte yok: {KATEGORI[indirim_kategori][0]}.")
    ind = yuvarla((ara[indirim_kategori] if indirim_kategori else toplam) * indirim / 100)
    return {"kalemler": sorted(kalemler, key=lambda x: list(KATEGORI).index(x["kategori"])), "ara": ara, "toplam": toplam, "indirim_orani": indirim, "indirim_kategori": indirim_kategori, "indirim": ind,
            "genel": toplam - ind, "para": paralar[0], "tarih": teklif_tarihi, "gecerli": teklif_tarihi + timedelta(days=gecerlilik), "bulunamayan": bulunamayan}


def teklif_metni(t: dict, dil: str = "tr", ref: str = "", notlar: tuple[str, ...] = ()) -> str:
    m, j, pb = METIN[dil], 0 if dil == "tr" else 1, t["para"]
    s = [m["baslik"], " · ".join(x for x in (f"{m['ref']}: {ref}" if ref else "", f"{m['tarih']}: {t['tarih']:%d.%m.%Y}", f"{m['gecerli']}: {t['gecerli']:%d.%m.%Y}") if x), ""]
    for k, a in t["ara"].items():
        s.append(KATEGORI[k][j])
        for x in t["kalemler"]:
            if x["kategori"] == k:
                s.append(f"  - {(x['en'] or x['ad']) if dil == 'en' else x['ad']} · {x['adet'].normalize():f} × {tutar(x['fiyat'], dil)} {pb} = {tutar(x['tutar'], dil)} {pb}".rstrip())
        s += [f"  {m['ara']}: {tutar(a, dil)} {pb}", ""]
    s.append(f"{m['toplam']}: {tutar(t['toplam'], dil)} {pb}")
    if t["indirim"]:
        kapsam = f", {KATEGORI[t['indirim_kategori']][j]}" if t["indirim_kategori"] else ""
        s.append(f"{m['indirim']} (%{t['indirim_orani'].normalize():f}{kapsam}): -{tutar(t['indirim'], dil)} {pb}")
    s += [f"{m['genel']}: {tutar(t['genel'], dil)} {pb}", "", f"{m['notlar']}:", f"- {m['n1']}", f"- {m['n2']}", *[f"- {n}" for n in notlar]]
    return "\n".join(s) + "\n"


def rapor_yaz(cikti: Path, t: dict, uy: list[dict], dil: str, ref: str) -> None:
    wb = Workbook()
    m, j = METIN[dil], 0 if dil == "tr" else 1
    ws = wb.active
    ws.title = "Teklif"
    ws.append([m["baslik"]])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"{m['ref']}: {ref}" if ref else "", "", f"{m['tarih']}: {t['tarih']:%d.%m.%Y}", "", f"{m['gecerli']}: {t['gecerli']:%d.%m.%Y}"])
    ws.append([])
    _baslik(ws, ["Kategori", "Kod", "Kalem", "Adet", "Birim Fiyat", "Tutar", "Para Birimi", "Not"], (16, 10, 52, 8, 14, 15, 10, 30), dondur="A5")
    for k, a in t["ara"].items():
        for x in t["kalemler"]:
            if x["kategori"] == k:
                _satir(ws, [KATEGORI[k][j], x["kod"], (x["en"] or x["ad"]) if dil == "en" else x["ad"], x["adet"], x["fiyat"], x["tutar"], t["para"], x["not"]], para_s=(5, 6), sar=(3, 8))
        n = _satir(ws, ["", "", f"{KATEGORI[k][j]} · {m['ara']}", None, None, a, t["para"]], para_s=(6,), dolgu={3: MAVI, 6: MAVI})
        ws.cell(n, 3).font = Font(bold=True)
    ws.append([])
    _satir(ws, ["", "", m["toplam"], None, None, t["toplam"], t["para"]], para_s=(6,))
    if t["indirim"]:
        _satir(ws, ["", "", f"{m['indirim']} (%{t['indirim_orani'].normalize():f}{', ' + KATEGORI[t['indirim_kategori']][j] if t['indirim_kategori'] else ''})", None, None, -t["indirim"], t["para"]], para_s=(6,))
    n = _satir(ws, ["", "", m["genel"], None, None, t["genel"], t["para"]], para_s=(6,), dolgu={3: YESIL, 6: YESIL})
    for c in ws[n]:
        c.font = Font(bold=True)
    ws.append([])
    ws.append([f"- {m['n1']}"])
    ws.append([f"- {m['n2']}"])
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(plan_yol: Path, fiyat_yol: Path, cikti: Path, teklif_tarihi: date, dil: str = "tr", ref: str = "", notlar: tuple[str, ...] = (), **ayar) -> dict:
    plan, liste, uy = oku(plan_yol, fiyat_yol)
    if not plan:
        raise ValueError("Tedavi planında kalem yok.")
    t = teklif(plan, liste, teklif_tarihi, **ayar)
    for b in t["bulunamayan"]:
        uy.append(uyari("Yüksek", "Kalem fiyat listesinde yok", b, "Teklife alınmadı. Fiyat listesindeki kodu veya adı aynen yazın; teklif bu kalem olmadan eksiktir."))
    if dil == "en":
        for x in t["kalemler"]:
            if not x["en"]:
                uy.append(uyari("Orta", "İngilizce adı yok", x["ad"], "Fiyat listesinde 'İngilizce Ad' boş; teklifte Türkçe adı kullanıldı."))
    if not t["para"]:
        uy.append(uyari("Orta", "Para birimi yok", fiyat_yol.name, "Fiyat listesinde para birimi boş; teklifte tutarların yanında para birimi yazmıyor."))
    t["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, t, uy, dil, ref)
    t["metin"] = teklif_metni(t, dil, ref, notlar)
    cikti.with_suffix(".txt").write_text(t["metin"], encoding="utf-8")
    return t


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Tedavi planı ve fiyat listesinden hasta teklifi hazırlar.")
    p.add_argument("--plan", type=Path, default=ORNEK / "tedavi_plani.csv", help="Kalem (fiyat listesindeki kod veya ad), Adet, Not")
    p.add_argument("--fiyatlar", type=Path, default=ORNEK / "fiyat_listesi.csv", help="Kod, Kalem, Kategori (Tedavi / Konaklama / Transfer / Diğer), Birim Fiyat, Para Birimi, İngilizce Ad")
    p.add_argument("--hasta-ref", help="Teklifte görünecek referans kodu (hasta adı yerine; örnek veride UH-2026-001)")
    p.add_argument("--indirim", help="İndirim %% (varsayılan 0; örnek veride 5)")
    p.add_argument("--indirim-kategori", choices=("tedavi", "konaklama", "transfer", "diger"), help="İndirimi yalnız bu kategoriye uygula (varsayılan: toplama)")
    p.add_argument("--gecerlilik-gun", default="30", help="Teklifin geçerlilik süresi (varsayılan 30 gün)")
    p.add_argument("--tarih", help="Teklif tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 12.10.2026)")
    p.add_argument("--dil", choices=("tr", "en"), default="tr", help="Teklif dili (varsayılan tr)")
    p.add_argument("--not", dest="notlar", action="append", default=[], help="Teklif metnine eklenecek not (birden çok kez verilebilir)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "tedavi_teklifi.xlsx")
    a = p.parse_args(argv)
    ornek = a.plan == ORNEK / "tedavi_plani.csv"
    tt = tarih(a.tarih) if a.tarih else (date(2026, 10, 12) if ornek else date.today())
    ind, gg = para(a.indirim or ("5" if ornek else "0")), para(a.gecerlilik_gun)
    if tt is None or ind is None or not 0 <= ind < 100 or gg is None or gg <= 0 or gg != int(gg):
        print("[X] --tarih GG.AA.YYYY biçiminde; --indirim 0 ile 100 arasında; --gecerlilik-gun pozitif tam sayı olmalı")
        return 2
    for y in (a.plan, a.fiyatlar):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        t = calistir(a.plan, a.fiyatlar, a.cikti, tt, a.dil, a.hasta_ref or ("UH-2026-001" if ornek else ""), tuple(a.notlar), indirim=ind, indirim_kategori=a.indirim_kategori, gecerlilik=int(gg))
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(t['kalemler'])} kalem · toplam {t['toplam']:,.2f} {t['para']} · indirim {t['indirim']:,.2f} · genel toplam {t['genel']:,.2f} {t['para']} · geçerlilik {t['gecerli']:%d.%m.%Y}")
    for b in t["bulunamayan"]:
        print(f"[!] Fiyat listesinde yok, teklife alınmadı: {b}")
    print(f"[OK] Teklif: {a.cikti.resolve()} ve {a.cikti.with_suffix('.txt').name} (göndermeden önce hekim ve kurum onayından geçirin)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
