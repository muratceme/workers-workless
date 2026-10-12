import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
TARIH = date(2026, 10, 12)


def liste(*kalemler, para="EUR"):
    l = {}
    for kod, ad, kat, fiyat in kalemler:
        x = {"kod": kod, "ad": ad, "kategori": kat, "fiyat": D(fiyat), "para": para, "en": f"{ad} EN" if kod != "X" else ""}
        l[kod.lower()] = l[ad.lower()] = x
    return l


L = liste(("t1", "implant", "tedavi", 600), ("t2", "kaplama", "tedavi", 180), ("k1", "otel", "konaklama", 70), ("r1", "transfer", "transfer", 40))


def pl(*kalemler):
    return [{"kalem": k, "adet": D(a), "not": ""} for k, a in kalemler]


class TeklifTesti(unittest.TestCase):
    def test_toplamlar_ve_indirim(self):
        t = main.teklif(pl(("T1", 4), ("kaplama", 8), ("k1", 5), ("r1", 2), ("yok", 1)), L, TARIH, D(5))
        self.assertEqual((t["ara"], t["toplam"], t["indirim"], t["genel"], t["para"], t["gecerli"], t["bulunamayan"]),
                         ({"tedavi": D("3840.00"), "konaklama": D("350.00"), "transfer": D("80.00")}, D(4270), D("213.50"), D("4056.50"), "EUR", date(2026, 11, 11), ["yok"]))
        t = main.teklif(pl(("r1", 2), ("t1", 1), ("k1", "2.5")), L, TARIH, D(10), "tedavi", 15)
        self.assertEqual(([x["kod"] for x in t["kalemler"]], t["indirim"], t["genel"], t["gecerli"]), (["t1", "k1", "r1"], D("60.00"), D("795.00"), date(2026, 10, 27)))
        self.assertEqual(main.teklif(pl(("t1", 1)), L, TARIH)["indirim"], D(0))

    def test_hatalar(self):
        for plan, lst, kw in ((pl(("yok", 1)), L, {}), (pl(("t1", 1), ("x", 1)), L | liste(("X", "usd kalem", "diger", 5), para="USD"), {}), (pl(("t1", 1)), L, {"indirim_kategori": "transfer"})):
            with self.assertRaises(ValueError):
                main.teklif(plan, lst, TARIH, **kw)

    def test_metin(self):
        t = main.teklif(pl(("t1", 4), ("k1", 5)), L, TARIH, D(5))
        tr = main.teklif_metni(t, "tr", "UH-1", ("Ulaşım uçak bileti hariçtir.",))
        self.assertIn("TEDAVİ FİYAT TEKLİFİ\nReferans: UH-1 · Teklif tarihi: 12.10.2026 · Geçerlilik: 11.11.2026\n\nTedavi\n  - implant · 4 × 600,00 EUR = 2.400,00 EUR\n  Ara toplam: 2.400,00 EUR\n", tr)
        self.assertIn("Toplam: 2.750,00 EUR\nİndirim (%5): -137,50 EUR\nGENEL TOPLAM: 2.612,50 EUR\n", tr)
        self.assertTrue(tr.endswith("- Teklifte yazmayan hizmetler fiyata dahil değildir.\n- Ulaşım uçak bileti hariçtir.\n"))
        en = main.teklif_metni(main.teklif(pl(("t1", 4), ("k1", 5)), L, TARIH, D(5), "konaklama"), "en")
        self.assertIn("TREATMENT PRICE QUOTE\nQuote date: 12.10.2026 · Valid until: 11.11.2026\n\nTreatment\n  - implant EN · 4 × 600.00 EUR = 2,400.00 EUR\n  Subtotal: 2,400.00 EUR\n\nAccommodation\n", en)
        self.assertIn("Total: 2,750.00 EUR\nDiscount (%5, Accommodation): -17.50 EUR\nGRAND TOTAL: 2,732.50 EUR\n", en)
        self.assertNotIn("İndirim", main.teklif_metni(main.teklif(pl(("t1", 1)), L, TARIH)))


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            t = main.calistir(ORNEK / "tedavi_plani.csv", ORNEK / "fiyat_listesi.csv", cikti, TARIH, ref="UH-2026-001", indirim=D(5))
            self.assertEqual((t["ara"], t["toplam"], t["indirim"], t["genel"], t["bulunamayan"], [u["tur"] for u in t["uyarilar"]]),
                             ({"tedavi": D("3870.00"), "konaklama": D("350.00"), "transfer": D("80.00")}, D(4300), D("215.00"), D("4085.00"), ["Tercüman hizmeti"], ["Kalem fiyat listesinde yok"]))
            self.assertEqual((Path(tmp) / "t.txt").read_text(encoding="utf-8"), t["metin"])
            self.assertIn("  - Zirkonyum kaplama (tek diş) · 8 × 180,00 EUR = 1.440,00 EUR", t["metin"])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Teklif", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            p, f = Path(tmp) / "p.csv", Path(tmp) / "f.csv"
            f.write_text("Kod;Kalem;Kategori;Birim Fiyat;Para Birimi\nA;Alfa;Tedavi;1.250,50;eur\nA;Alfa iki;Bilinmeyen;5;EUR\nB;Beta;Konaklama;pahalı;EUR\n;Gama;;10;EUR\n", encoding="utf-8")
            p.write_text("Kalem;Adet\nA;2\nGAMA;\nalfa;sıfır\n;3\n", encoding="utf-8")
            plan, lst, uy = main.oku(p, f)
            self.assertEqual(([(x["kalem"], x["adet"]) for x in plan], {k: (v["ad"], v["kategori"], v["fiyat"], v["para"]) for k, v in lst.items()}),
                             ([("A", D(2)), ("GAMA", D(1))], {"a": ("Alfa", "tedavi", D("1250.50"), "EUR"), "alfa": ("Alfa", "tedavi", D("1250.50"), "EUR"), "alfa iki": ("Alfa iki", "diger", D(5), "EUR"),
                                                                 "gama": ("Gama", "diger", D(10), "EUR")}))
            self.assertEqual(sorted(u["tur"] for u in uy), ["Adet okunamadı", "Fiyat listesinde mükerrer kalem", "Liste fiyatı okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--plan", str(p), "--fiyatlar", str(f), "--cikti", c, "--tarih", "12.10.2026"]
            self.assertEqual(main.main(arg + ["--dil", "en", "--indirim", "7,5", "--indirim-kategori", "tedavi", "--gecerlilik-gun", "15", "--hasta-ref", "R-1", "--not", "Flights are not included."]), 0)
            self.assertIn("GRAND TOTAL: 2,323.42 EUR", (Path(tmp) / "x.txt").read_text(encoding="utf-8"))            # 2.511,00 − (2.501,00 × %7,5 = 187,58)
            self.assertEqual(main.main(arg + ["--indirim", "100"]), 2)
            self.assertEqual(main.main(arg + ["--gecerlilik-gun", "0"]), 2)
            self.assertEqual(main.main(arg + ["--indirim-kategori", "transfer"]), 1)
            self.assertEqual(main.main(["--plan", str(Path(tmp) / "yok.csv")]), 1)
            p.write_text("Kalem;Adet\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
