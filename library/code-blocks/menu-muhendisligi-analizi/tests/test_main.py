import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def k(urun, adet, fiyat, maliyet, kategori="Ana"):
    return {"urun": urun, "kategori": kategori, "adet": D(adet), "fiyat": D(fiyat), "maliyet": D(maliyet)}


ANA = lambda: [k("A", 400, 500, 150), k("B", 300, 400, 200), k("C", 100, 600, 180), k("D", 200, 300, 150), k("E", 50, 250, 150)]   # noqa: E731


class AnalizTesti(unittest.TestCase):
    def test_siniflandirma(self):
        r = main.analiz(ANA())
        self.assertEqual([(x["urun"], x["sinif"], x["karisim"], x["katki"], x["toplam_katki"], x["maliyet_y"], x["katki_pay"]) for x in r["kalemler"]],
                         [("A", "Yıldız", D("38.10"), D("350.00"), D("140000.00"), D("30.00"), D("50.54")), ("B", "At", D("28.57"), D("200.00"), D("60000.00"), D("50.00"), D("21.66")),
                          ("D", "At", D("19.05"), D("150.00"), D("30000.00"), D("50.00"), D("10.83")), ("C", "Bulmaca", D("9.52"), D("420.00"), D("42000.00"), D("30.00"), D("15.16")),
                          ("E", "Köpek", D("4.76"), D("100.00"), D("5000.00"), D("60.00"), D("1.81"))])
        o = r["kategoriler"][0]
        self.assertEqual((o["kalem"], o["adet"], o["ciro"], o["toplam_katki"], o["ort_katki"], o["esik"], o["maliyet_y"], o["Yıldız"], o["At"], o["Bulmaca"], o["Köpek"], r["say"], r["uyarilar"]),
                         (5, D(1050), D("452500.00"), D("277000.00"), D("263.81"), D("14.00"), D("38.78"), 1, 2, 1, 1, {"Yıldız": 1, "At": 2, "Bulmaca": 1, "Köpek": 1}, []))
        self.assertIn("Menüden çıkarmayı", r["kalemler"][-1]["aksiyon"])

    def test_esik_kdv_ve_sinirlar(self):
        s = lambda l, **kw: {x["urun"]: x["sinif"] for x in main.analiz(l, **kw)["kalemler"]}   # noqa: E731
        self.assertEqual((s(ANA(), esik=D(100))["D"], s(ANA(), esik=D(20))["E"], s(ANA(), esik=D(20))["C"]), ("Köpek", "At", "Yıldız"))          # eşik %20 → %4; %100 → %20
        iki = [k("X", 50, 100, 40), k("Y", 50, 100, 40)]
        self.assertEqual(s(iki), {"X": "Yıldız", "Y": "Yıldız"})                                 # eşitlik: eşik ve ortalamaya eşit olan üstte sayılır
        self.assertEqual(s([k("X", 35, 100, 40), k("Y", 65, 100, 41)]), {"X": "Yıldız", "Y": "At"})
        self.assertEqual(s([k("X", 34, 100, 40), k("Y", 66, 100, 41)]), {"X": "Bulmaca", "Y": "At"})
        x = main.analiz(ANA(), kdv=D(10))["kalemler"][0]
        self.assertEqual((x["urun"], x["net"], x["katki"], x["maliyet_y"]), ("A", D("454.55"), D("304.55"), D("33.00")))

    def test_kategoriler_ve_uyarilar(self):
        l = ANA() + [k("T1", 10, 100, 120, "Tatlı"), k("İ1", 0, 50, 10, "İçecek"), k("İ2", 0, 60, 10, "içecek")]
        r = main.analiz(l)
        self.assertEqual(([o["kategori"] for o in r["kategoriler"]], [x["sinif"] for x in r["kalemler"] if x["kategori"] == "Tatlı"], sorted((u["tur"], u["kim"]) for u in r["uyarilar"])),
                         (["Ana", "Tatlı"], ["Yıldız"], [("Maliyet satış fiyatından yüksek", "T1"), ("Satışı olmayan kategori", "İçecek"), ("Tek kalemli kategori", "Tatlı")]))
        r = main.analiz(l, tum_menu=True)
        self.assertEqual(([(o["kategori"], o["kalem"], o["adet"], o["esik"]) for o in r["kategoriler"]], {x["urun"]: x["sinif"] for x in r["kalemler"]}["T1"], len(r["kalemler"])), ([("Tüm Menü", 8, D(1060), D("8.75"))], "Köpek", 8))
        for kalemler in ([], [k("X", 0, 10, 1)]):
            with self.assertRaises(ValueError):
                main.analiz(kalemler)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "m.xlsx"
            r = main.calistir(ORNEK / "menu_satislari.csv", cikti)
            self.assertEqual([(x["urun"], x["sinif"]) for x in r["kalemler"]], [("Izgara Köfte", "Yıldız"), ("Tavuk Şiş", "At"), ("Sebzeli Makarna", "At"), ("Kuzu Tandır", "Bulmaca"), ("Balık Buğulama", "Köpek"),
                                                                              ("Sütlaç", "Yıldız"), ("Künefe", "Yıldız"), ("Meyve Tabağı", "Köpek")])
            t = r["kategoriler"][1]
            self.assertEqual((t["kategori"], t["adet"], t["toplam_katki"], t["ort_katki"], t["esik"], r["say"], r["uyarilar"]), ("Tatlı", D(500), D("74500.00"), D("149.00"), D("23.33"), {"Yıldız": 3, "At": 2, "Bulmaca": 1, "Köpek": 2}, []))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Analiz", "Kategori Özeti", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "m.csv"
            y.write_text("Ürün;Adet;Fiyat;Maliyet\nA;1.250;350,50;120\nB;x;10;5\nC;5;0;5\nD;5;10;-1\nE;0;10;5\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["urun"], x["kategori"], x["adet"], x["fiyat"]) for x in l], [u["tur"] for u in uy]), ([("A", "Genel", D(1250), D("350.50")), ("E", "Genel", D(0), D(10))], ["Kalem okunamadı"] * 3))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--satislar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--esik", "60", "--kdv", "10", "--tum-menu"]), 0)
            for ek in (["--esik", "0"], ["--kdv", "100"], ["--kdv", "x"]):
                self.assertEqual(main.main(arg + ek), 2)
            self.assertEqual(main.main(["--satislar", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Ürün;Adet\nA;1\n", "Ürün;Adet;Fiyat;Maliyet\nA;0;10;5\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
