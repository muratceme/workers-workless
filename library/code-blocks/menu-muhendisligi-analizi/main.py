"""
Menü Mühendisliği Analizi — Gerisi Bende kod bloğu
Yiyecek ve İçecek › Yiyecek İçecek Müdürü

Satış adetleri ve katkı paylarından menü kalemlerini yıldız, at, bulmaca ve köpek olarak sınıflandırır (Kasavana-Smith yöntemi), aksiyon önerir.
  Katkı payı          = satış fiyatı (KDV hariç) − porsiyon maliyeti;   toplam katkı = katkı payı × satış adedi
  Menü karışımı %     = kalemin satış adedi ÷ kategorinin toplam adedi × 100
  Popülerlik eşiği %  = (100 ÷ kategorideki kalem sayısı) × --esik ÷ 100                    (varsayılan 70: "beklenen payın %70'i")
  Ortalama katkı      = kategorinin toplam katkısı ÷ toplam adedi                           (ağırlıklı ortalama)
  Sınıf               = Yıldız (popüler, katkısı yüksek) · At (popüler, katkısı düşük) · Bulmaca (popüler değil, katkısı yüksek) · Köpek (ikisi de düşük)
                        popüler: menü karışımı ≥ eşik;  katkısı yüksek: katkı payı ≥ ortalama katkı
  Maliyet oranı %     = porsiyon maliyeti ÷ satış fiyatı (KDV hariç) × 100
Kalemler kendi kategorisi içinde karşılaştırılır (ana yemek ile tatlı kıyaslanmaz); --tum-menu ile tek grup yapılır. Fiyatlar KDV dahilse --kdv oranını verin.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"urun": ("urun", "menu kalemi", "kalem", "urun adi", "yemek"), "kategori": ("kategori", "grup", "menu grubu"), "adet": ("satis adedi", "adet", "satilan adet", "satis"),
         "fiyat": ("satis fiyati", "fiyat"), "maliyet": ("porsiyon maliyeti", "maliyet", "birim maliyet")}
Z = Decimal(0)
SINIF = {(True, True): "Yıldız", (True, False): "At", (False, True): "Bulmaca", (False, False): "Köpek"}
AKSIYON = {"Yıldız": "Koruyun: kaliteyi, porsiyonu ve menüdeki görünür yerini sabit tutun. Fiyat artışına en dayanıklı kalemdir.",
           "At": "Kârlılığı artırın: küçük fiyat artışı deneyin, reçete ve porsiyon maliyetini gözden geçirin, katkısı yüksek yan ürünle birlikte sunun.",
           "Bulmaca": "Satışı artırın: menüde öne alın, servis ekibine önerdirin, adını ve tanımını gözden geçirin. Fiyat yüksek algılanıyorsa test edin.",
           "Köpek": "Menüden çıkarmayı veya yeniden tasarlamayı değerlendirin. Tutulacaksa öne çıkarmayın, fiyatını gözden geçirin."}
SINIF_RENK = {"Yıldız": YESIL, "At": MAVI, "Bulmaca": SARI, "Köpek": KIRMIZI}


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("urun", "adet", "fiyat", "maliyet")):
        u = metin(r["urun"])
        if not u:
            continue
        a, f, m = para(r["adet"]), para(r["fiyat"]), para(r["maliyet"])
        if a is None or a < 0 or f is None or f <= 0 or m is None or m < 0:
            uy.append(uyari("Yüksek", "Kalem okunamadı", f"{u} (satır {r['_satir']})", f"Adet '{metin(r['adet'])}', fiyat '{metin(r['fiyat'])}', maliyet '{metin(r['maliyet'])}'. Adet ve maliyet sıfır veya pozitif, fiyat pozitif olmalı. Satır alınmadı."))
            continue
        l.append({"urun": u, "kategori": metin(r["kategori"]) or "Genel", "adet": a, "fiyat": f, "maliyet": m})
    return l, uy


def analiz(kalemler: list[dict], esik: Decimal = Decimal(70), kdv: Decimal = Z, tum_menu: bool = False) -> dict:
    if not kalemler:
        raise ValueError("Analiz edilecek menü kalemi yok.")
    uy, grup = [], {}
    for k in kalemler:
        ad = "Tüm Menü" if tum_menu else k["kategori"]
        grup.setdefault(katla(ad), {"kategori": ad, "kalemler": []})["kalemler"].append(k)
    liste, ozet = [], []
    for g in grup.values():
        ks, n = g["kalemler"], len(g["kalemler"])
        hazir = [k | {"net": k["fiyat"] / (1 + kdv / 100)} for k in ks]
        for k in hazir:
            k["katki"] = k["net"] - k["maliyet"]
            if k["katki"] < 0:
                uy.append(uyari("Yüksek", "Maliyet satış fiyatından yüksek", k["urun"], f"KDV hariç fiyat {yuvarla(k['net']):,.2f}, porsiyon maliyeti {k['maliyet']:,.2f}. Her satışta zarar ediliyor; fiyatı veya reçeteyi kontrol edin."))
        N = sum(k["adet"] for k in hazir)
        if N == 0:
            uy.append(uyari("Orta", "Satışı olmayan kategori", g["kategori"], "Kategorideki hiçbir kalemin satış adedi yok; sınıflandırma yapılamadı."))
            continue
        if n == 1:
            uy.append(uyari("Bilgi", "Tek kalemli kategori", g["kategori"], "Karşılaştırılacak başka kalem yok; kalem kendi ortalamasıyla kıyaslandığı için 'Yıldız' çıkar. Kategorileri birleştirmeyi veya --tum-menu seçeneğini düşünün."))
        top_katki, ciro, top_maliyet = sum(k["katki"] * k["adet"] for k in hazir), sum(k["net"] * k["adet"] for k in hazir), sum(k["maliyet"] * k["adet"] for k in hazir)
        ort, pop_esik = top_katki / N, Decimal(100) / n * esik / 100
        say = defaultdict(int)
        for k in hazir:
            mm = k["adet"] / N * 100
            sinif = SINIF[(mm >= pop_esik, k["katki"] >= ort)]
            say[sinif] += 1
            liste.append({"urun": k["urun"], "kategori": g["kategori"], "adet": k["adet"], "fiyat": k["fiyat"], "net": yuvarla(k["net"]), "maliyet": k["maliyet"], "maliyet_y": yuvarla(k["maliyet"] / k["net"] * 100),
                          "katki": yuvarla(k["katki"]), "toplam_katki": yuvarla(k["katki"] * k["adet"]), "karisim": yuvarla(mm), "katki_pay": yuvarla(k["katki"] * k["adet"] / top_katki * 100) if top_katki else None,
                          "populer": mm >= pop_esik, "karli": k["katki"] >= ort, "sinif": sinif, "aksiyon": AKSIYON[sinif]})
        ozet.append({"kategori": g["kategori"], "kalem": n, "adet": N, "ciro": yuvarla(ciro), "toplam_katki": yuvarla(top_katki), "ort_katki": yuvarla(ort), "esik": yuvarla(pop_esik), "maliyet_y": yuvarla(top_maliyet / ciro * 100)}
                    | {s: say[s] for s in ("Yıldız", "At", "Bulmaca", "Köpek")})
    if not liste:
        raise ValueError("Hiçbir kategoride satış yok; sınıflandırma yapılamadı.")
    sira = {"Yıldız": 0, "At": 1, "Bulmaca": 2, "Köpek": 3}
    liste.sort(key=lambda x: (x["kategori"], sira[x["sinif"]], -x["toplam_katki"], x["urun"]))
    return {"kalemler": liste, "kategoriler": sorted(ozet, key=lambda o: o["kategori"]), "say": {s: sum(1 for x in liste if x["sinif"] == s) for s in sira}, "ayar": (esik, kdv), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    esik, kdv = r["ayar"]
    ws = wb.active
    ws.title = "Analiz"
    _baslik(ws, ["Kategori", "Menü Kalemi", "Sınıf", "Satış Adedi", "Menü Karışımı %", "Satış Fiyatı", "KDV Hariç Fiyat", "Porsiyon Maliyeti", "Maliyet Oranı %", "Katkı Payı", "Toplam Katkı", "Katkıdaki Payı %", "Önerilen Aksiyon",
                 "Karar (siz doldurun)"], (18, 30, 10, 9, 10, 12, 12, 12, 10, 12, 14, 10, 80, 24), "C2")
    for x in r["kalemler"]:
        _satir(ws, [x["kategori"], x["urun"], x["sinif"], x["adet"], x["karisim"], x["fiyat"], x["net"], x["maliyet"], x["maliyet_y"], x["katki"], x["toplam_katki"], x["katki_pay"], x["aksiyon"], ""],
               para_s=(5, 6, 7, 8, 9, 10, 11, 12), dolgu={3: SINIF_RENK[x["sinif"]], 5: None if x["populer"] else SARI, 10: None if x["karli"] else SARI}, sar=(13,))
    ws.append([])
    ws.append([f"Popüler: menü karışımı ≥ (100 ÷ kalem sayısı) × %{esik.normalize():f}. Katkısı yüksek: katkı payı ≥ kategorinin ağırlıklı ortalama katkısı. Sarı hücre eşiğin altında kalan ölçüttür."
               + (f" Fiyatlardan %{kdv.normalize():f} KDV ayrıştırıldı." if kdv else " Fiyatlar KDV hariç kabul edildi.")])
    ws = wb.create_sheet("Kategori Özeti")
    _baslik(ws, ["Kategori", "Kalem Sayısı", "Satış Adedi", "Ciro (KDV hariç)", "Toplam Katkı", "Ortalama Katkı", "Popülerlik Eşiği %", "Maliyet Oranı %", "Yıldız", "At", "Bulmaca", "Köpek"], (20, 9, 10, 16, 15, 12, 11, 10, 8, 8, 9, 8))
    for o in r["kategoriler"]:
        _satir(ws, [o["kategori"], o["kalem"], o["adet"], o["ciro"], o["toplam_katki"], o["ort_katki"], o["esik"], o["maliyet_y"], o["Yıldız"] or None, o["At"] or None, o["Bulmaca"] or None, o["Köpek"] or None], para_s=(4, 5, 6, 7, 8))
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, **ayar) -> dict:
    kalemler, uy = oku(yol)
    r = analiz(kalemler, **ayar)
    uy += r["uyarilar"]
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Menü kalemlerini satış adedi ve katkı payına göre yıldız, at, bulmaca, köpek olarak sınıflandırır.")
    p.add_argument("--satislar", type=Path, default=ORNEK / "menu_satislari.csv", help="Menü Kalemi, Kategori, Satış Adedi, Satış Fiyatı, Porsiyon Maliyeti")
    p.add_argument("--esik", default="70", help="Popülerlik eşiği: beklenen payın yüzdesi (varsayılan 70)")
    p.add_argument("--kdv", default="0", help="Satış fiyatları KDV dahilse oran, %% (varsayılan 0: fiyatlar KDV hariç)")
    p.add_argument("--tum-menu", action="store_true", help="Kategorileri ayırmadan bütün menüyü tek grup olarak değerlendir")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "menu_muhendisligi.xlsx")
    a = p.parse_args(argv)
    e, k = para(a.esik), para(a.kdv)
    if e is None or not 0 < e <= 100 or k is None or not 0 <= k < 100:
        print("[X] --esik 0–100 arasında, --kdv 0 ile 100 arasında bir sayı olmalı")
        return 2
    if not a.satislar.exists():
        print(f"[X] Dosya bulunamadı: {a.satislar}")
        return 1
    try:
        r = calistir(a.satislar, a.cikti, esik=e, kdv=k, tum_menu=a.tum_menu)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    s = r["say"]
    print(f"[OK] {len(r['kalemler'])} kalem · {len(r['kategoriler'])} kategori · {s['Yıldız']} yıldız · {s['At']} at · {s['Bulmaca']} bulmaca · {s['Köpek']} köpek")
    for x in r["kalemler"]:
        if x["sinif"] == "Köpek":
            print(f"[!] Köpek: {x['urun']} ({x['kategori']}) · menü karışımı %{x['karisim']}, katkı payı {x['katki']:,.2f}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
