import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)
KUR = {"TL": D(1), "USD": D(40)}


def s(banka, limit, kullanim, sinif="Nakdi", pb="TL", faiz=None, bitis=None, teminat=None):
    return {"banka": banka, "urun": "Ü", "sinif": sinif, "pb": pb, "limit": D(limit), "kullanim": D(kullanim), "faiz": None if faiz is None else D(faiz), "bitis": bitis, "teminat": "", "teminat_tutar": None if teminat is None else D(teminat),
            "kim": f"{banka} · Ü"}


class RaporTesti(unittest.TestCase):
    def test_satir_ve_banka(self):
        r = main.raporla([s("A", 1000, 600, faiz=40, teminat=900), s("A", 500, 100, "Gayrinakdi"), s("B", 100, 50, pb="USD", faiz=8), s("B", 1000, 200, faiz=50), s("C", 100, 10, pb="EUR", faiz=5)], KUR, BUGUN)
        self.assertEqual([(x["bos"], x["oran"], x["tl_limit"], x["tl_kullanim"], x["teminat_oran"]) for x in r["satirlar"]],
                         [(D(400), D("60.00"), D("1000.00"), D("600.00"), D("150.00")), (D(400), D("20.00"), D("500.00"), D("100.00"), None), (D(50), D("50.00"), D("4000.00"), D("2000.00"), None),
                          (D(800), D("20.00"), D("1000.00"), D("200.00"), None), (D(90), D("10.00"), None, None, None)])
        self.assertEqual([(b["banka"], b["nakdi_limit"], b["nakdi"], b["gn_limit"], b["gn"], b["risk"], b["bos"], b["pay"]) for b in r["bankalar"]],
                         [("B", D(5000), D(2200), D(0), D(0), D(2200), D(2800), D("75.86")), ("A", D(1000), D(600), D(500), D(100), D(700), D(800), D("24.14"))])
        self.assertEqual((r["toplam"], r["faiz"]), ({"nakdi_limit": D(6000), "nakdi": D(2800), "gn_limit": D(500), "gn": D(100), "risk": D(2900), "bos": D(3600)}, {"EUR": D("5.00"), "TL": D("42.50"), "USD": D("8.00")}))
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]), [("Banka yoğunlaşması", "B"), ("Kur bulunamadı", "C · Ü")])

    def test_uyarilar(self):
        l = [s("A", 100, 120, faiz=40), s("B", 100, 90), s("C", 100, 89, "Gayrinakdi", bitis=date(2026, 10, 11)), s("D", 100, 0, bitis=date(2026, 11, 11)), s("E", 100, 0, bitis=date(2026, 11, 12))]
        r = main.raporla(l, KUR, BUGUN)
        self.assertEqual([(x["bos"], x["oran"], x["kalan_gun"]) for x in r["satirlar"]], [(D(-20), D("120.00"), None), (D(10), D("90.00"), None), (D(11), D("89.00"), -1), (D(100), D("0.00"), 30), (D(100), D("0.00"), 31)])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Bilgi", "Banka yoğunlaşması", "A · Ü"[:1]), ("Bilgi", "Faiz oranı yok", "B · Ü"), ("Orta", "Limit dolmak üzere", "B · Ü"), ("Orta", "Limit süresi doluyor", "D · Ü"), ("Yüksek", "Limit aşımı", "A · Ü"),
                          ("Yüksek", "Limit süresi dolmuş", "C · Ü")])
        r = main.raporla([s("B", 100, 90), s("D", 100, 0, bitis=date(2026, 11, 12))], KUR, BUGUN, D(95), 31, D(100))
        self.assertEqual(sorted(u["tur"] for u in r["uyarilar"]), ["Banka yoğunlaşması", "Faiz oranı yok", "Limit süresi doluyor"])
        self.assertEqual([u["tur"] for u in main.raporla([s("A", 100, 50, faiz=1)], KUR, BUGUN)["uyarilar"]], [])
        with self.assertRaises(ValueError):
            main.raporla([], KUR, BUGUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "b.xlsx"
            r = main.calistir(ORNEK / "limitler.csv", cikti, BUGUN, ORNEK / "kurlar.csv")
            self.assertEqual([(b["banka"], b["nakdi"], b["gn"], b["risk"], b["pay"]) for b in r["bankalar"]],
                             [("B Bankası", D(14600000), D(0), D(14600000), D("58.40")), ("A Bankası", D(4600000), D(1250000), D(5850000), D("23.40")), ("C Bankası", D(2100000), D(2450000), D(4550000), D("18.20")),
                              ("D Bankası", D(0), D(0), D(0), D("0.00"))])
            self.assertEqual((r["toplam"], r["faiz"]),
                             ({"nakdi_limit": D(32500000), "nakdi": D(21300000), "gn_limit": D(12800000), "gn": D(3700000), "risk": D(25000000), "bos": D(20300000)}, {"TL": D("50.60"), "USD": D("8.50")}))
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Yüksek", "Limit aşımı", "C Bankası · Taksitli Ticari Kredi"), ("Yüksek", "Limit süresi dolmuş", "C Bankası · Teminat Mektubu"), ("Orta", "Limit dolmak üzere", "A Bankası · Rotatif Kredi"),
                              ("Orta", "Limit süresi doluyor", "B Bankası · Spot Kredi"), ("Bilgi", "Banka yoğunlaşması", "B Bankası")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Limit ve Risk", "Banka Özeti", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, k = Path(tmp) / "l.csv", Path(tmp) / "k.csv"
            y.write_text("Banka Adı;Kredi Türü;Nakdi/Gayrinakdi;Döviz;Limit Tutarı;Risk;Faiz\nA;Rotatif;;;1.000.000,00;250.000;45,5\nA;Mektup;Gayri Nakdi;usd;100.000;0;\nB;Spot;Nakdi;TRY;0;10;\nC;Spot;Nakdi;TL;100;-5;\n"
                         "D;Spot;Nakdi;TL;x;1;\n", encoding="utf-8")
            k.write_text("Döviz;Kur\nusd;40,5\nEUR;x\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["banka"], x["urun"], x["sinif"], x["pb"], x["limit"], x["kullanim"], x["faiz"]) for x in l], [u["tur"] for u in uy]),
                             ([("A", "Rotatif", "Nakdi", "TL", D("1000000.00"), D(250000), D("45.5")), ("A", "Mektup", "Gayrinakdi", "USD", D(100000), D(0), None)], ["Sınıf yazılmamış"] + ["Satır okunamadı"] * 3))
            self.assertEqual(main.kur_oku(k), {"TL": D(1), "USD": D("40.5")})
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--limitler", str(y), "--cikti", c]
            self.assertEqual(main.calistir(y, Path(c), BUGUN, k)["toplam"]["gn_limit"], D("4050000.00"))
            self.assertEqual(main.main(arg + ["--kurlar", str(k), "--tarih", "12.10.2026", "--esik", "85", "--gun", "45", "--yogunlasma", "50"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual((main.main(arg + ["--esik", "0"]), main.main(arg + ["--yogunlasma", "x"]), main.main(arg + ["--gun", "-1"]), main.main(arg + ["--tarih", "x"])), (2, 2, 2, 2))
            self.assertEqual((main.main(["--limitler", str(Path(tmp) / "yok.csv")]), main.main(arg + ["--kurlar", str(Path(tmp) / "yok.csv")])), (1, 1))
            for icerik in ("Banka;Limit\nA;1\n", "Banka;Limit;Kullanım\nA;x;1\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
