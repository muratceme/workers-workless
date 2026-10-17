"""
Banka Limit ve Risk Raporu — Gerisi Bende kod bloğu
Finans › Finans Uzmanı

Banka bazında kredi limitlerini, kullanımları, teminatları ve faiz oranlarını tek tabloda toplar.
  Boş limit           = limit − kullanım;   kullanım oranı = kullanım ÷ limit
  TL karşılığı        = tutar × kur (--kurlar dosyasından; TL = 1). Kuru olmayan para birimi TL toplamlarına girmez
  Toplam risk         = nakdi kullanım + gayrinakdi kullanım (TL)
  Banka payı          = bankanın toplam riski ÷ tüm bankaların toplam riski
  Ağırlıklı faiz      = Σ(kullanım × faiz) ÷ Σ kullanım — yalnız nakdi, kullanımı ve faizi olan satırlar; para birimi bazında
  Uyarılar            = limit aşımı, kullanım oranı ≥ --esik (%90), süresi dolmuş veya --gun (30) gün içinde dolacak limit,
                        payı ≥ --yogunlasma (%40) olan banka
Rapor sizin kayıtlarınıza dayanır; bankanın sistemindeki limit ve riskle mutabakat yerine geçmez.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
SUTUN = {"banka": ("banka", "banka adi"), "urun": ("urun", "kredi turu", "limit turu", "tur"), "sinif": ("sinif", "nakdi gayrinakdi", "risk sinifi"), "pb": ("para birimi", "doviz", "doviz cinsi", "pb"),
         "limit": ("limit", "limit tutari"), "kullanim": ("kullanim", "risk", "kullanilan", "bakiye"), "faiz": ("faiz orani", "faiz", "oran"), "bitis": ("limit bitis", "bitis", "vade", "revize tarihi", "limit vadesi"),
         "teminat": ("teminat", "teminat turu"), "teminat_tutar": ("teminat tutari",)}
KUR_SUTUN = {"pb": ("para birimi", "doviz", "doviz cinsi", "pb"), "kur": ("kur", "doviz kuru")}
Z, YUZ = Decimal(0), Decimal(100)


def pb_duzelt(x) -> str:
    s = metin(x).upper().replace("₺", "TL")
    return "TL" if s in ("", "TL", "TRY", "TRL") else s


def kur_oku(yol: Path) -> dict[str, Decimal]:
    k = {"TL": Decimal(1)}
    for r in kayitlar(yol, KUR_SUTUN, ("pb", "kur")):
        v = para(r["kur"])
        if metin(r["pb"]) and v and v > 0:
            k[pb_duzelt(r["pb"])] = v
    return k


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("banka", "limit", "kullanim")):
        b = metin(r["banka"])
        if not b:
            continue
        kim = f"{b} · {metin(r['urun'])}".strip(" ·")
        lim, kul = para(r["limit"]), para(r["kullanim"])
        if lim is None or lim <= 0 or kul is None or kul < 0:
            uy.append(uyari("Yüksek", "Satır okunamadı", f"{kim} (satır {r['_satir']})", f"Limit '{metin(r['limit'])}', kullanım '{metin(r['kullanim'])}'. Limit pozitif, kullanım sıfır veya pozitif olmalı. Satır rapora alınmadı."))
            continue
        s = katla(r["sinif"])
        if not s:
            uy.append(uyari("Bilgi", "Sınıf yazılmamış", kim, "Nakdi / Gayrinakdi sütunu boş; satır nakdi sayıldı."))
        l.append({"banka": b, "urun": metin(r["urun"]), "sinif": "Gayrinakdi" if s.startswith("gayri") or s == "gn" else "Nakdi", "pb": pb_duzelt(r["pb"]), "limit": lim, "kullanim": kul, "faiz": para(r["faiz"]),
                  "bitis": tarih(r["bitis"]), "teminat": metin(r["teminat"]), "teminat_tutar": para(r["teminat_tutar"]), "kim": kim})
    return l, uy


def raporla(satirlar: list[dict], kurlar: dict[str, Decimal], bugun: date, esik: Decimal = Decimal(90), gun: int = 30, yogunlasma: Decimal = Decimal(40)) -> dict:
    if not satirlar:
        raise ValueError("Raporlanacak limit satırı yok.")
    uy, bankalar, faiz = [], {}, {}
    for s in satirlar:
        kur = kurlar.get(s["pb"])
        s["bos"], s["oran"] = s["limit"] - s["kullanim"], yuvarla(s["kullanim"] * YUZ / s["limit"])
        s["tl_limit"], s["tl_kullanim"] = (yuvarla(s["limit"] * kur), yuvarla(s["kullanim"] * kur)) if kur else (None, None)
        s["kalan_gun"] = (s["bitis"] - bugun).days if s["bitis"] else None
        s["teminat_oran"] = yuvarla(s["teminat_tutar"] * YUZ / s["kullanim"]) if s["teminat_tutar"] and s["kullanim"] else None
        if kur is None:
            uy.append(uyari("Yüksek", "Kur bulunamadı", s["kim"], f"{s['pb']} için kur verilmedi; satır TL toplamlarına ve banka paylarına girmedi. --kurlar dosyasına ekleyin."))
        if s["bos"] < 0:
            uy.append(uyari("Yüksek", "Limit aşımı", s["kim"], f"Kullanım {s['kullanim']:,.2f} {s['pb']}, limit {s['limit']:,.2f} {s['pb']}; aşım {-s['bos']:,.2f} {s['pb']} (%{s['oran']})."))
        elif s["oran"] >= esik:
            uy.append(uyari("Orta", "Limit dolmak üzere", s["kim"], f"Kullanım oranı %{s['oran']}; boş limit {s['bos']:,.2f} {s['pb']}."))
        if s["kalan_gun"] is not None and s["kalan_gun"] < 0:
            uy.append(uyari("Yüksek", "Limit süresi dolmuş", s["kim"], f"Limit bitişi {s['bitis']:%d.%m.%Y}, {-s['kalan_gun']} gün geçmiş. Yenileme durumunu bankayla doğrulayın."))
        elif s["kalan_gun"] is not None and s["kalan_gun"] <= gun:
            uy.append(uyari("Orta", "Limit süresi doluyor", s["kim"], f"Limit bitişi {s['bitis']:%d.%m.%Y}, {s['kalan_gun']} gün kaldı."))
        if s["sinif"] == "Nakdi" and s["kullanim"] > 0:
            if s["faiz"] is None:
                uy.append(uyari("Bilgi", "Faiz oranı yok", s["kim"], "Kullanımı olan nakdi kredide faiz oranı yazılmamış; ağırlıklı ortalamaya girmedi."))
            else:
                f = faiz.setdefault(s["pb"], [Z, Z])
                f[0] += s["kullanim"] * s["faiz"]
                f[1] += s["kullanim"]
        if kur:
            b = bankalar.setdefault(katla(s["banka"]), {"banka": s["banka"], "nakdi_limit": Z, "nakdi": Z, "gn_limit": Z, "gn": Z})
            on = "nakdi" if s["sinif"] == "Nakdi" else "gn"
            b[on + "_limit"] += s["tl_limit"]
            b[on] += s["tl_kullanim"]
    top = {k: sum((b[k] for b in bankalar.values()), Z) for k in ("nakdi_limit", "nakdi", "gn_limit", "gn")}
    top["risk"], top["bos"] = top["nakdi"] + top["gn"], top["nakdi_limit"] + top["gn_limit"] - top["nakdi"] - top["gn"]
    ozet = []
    for b in bankalar.values():
        risk = b["nakdi"] + b["gn"]
        pay = yuvarla(risk * YUZ / top["risk"]) if top["risk"] else Z
        ozet.append(b | {"risk": risk, "bos": b["nakdi_limit"] + b["gn_limit"] - risk, "pay": pay})
        if pay >= yogunlasma and len(bankalar) > 1:
            uy.append(uyari("Bilgi", "Banka yoğunlaşması", b["banka"], f"Toplam riskin %{pay} kadarı bu bankada. Tek bankaya bağımlılığı değerlendirin."))
    return {"satirlar": satirlar, "bankalar": sorted(ozet, key=lambda b: (-b["risk"], katla(b["banka"]))), "toplam": top,
            "faiz": {pb: yuvarla(f[0] / f[1]) for pb, f in sorted(faiz.items())}, "bugun": bugun, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Limit ve Risk"
    _baslik(ws, ["Banka", "Ürün", "Sınıf", "Para Birimi", "Limit", "Kullanım", "Boş Limit", "Kullanım %", "Faiz %", "Limit Bitiş", "Kalan Gün", "Teminat", "Teminat Tutarı", "Teminat / Risk %", "Limit (TL)", "Kullanım (TL)"],
            (20, 26, 11, 8, 16, 16, 16, 10, 8, 12, 9, 26, 16, 10, 16, 16))
    for s in r["satirlar"]:
        _satir(ws, [s["banka"], s["urun"], s["sinif"], s["pb"], s["limit"], s["kullanim"], s["bos"], s["oran"], s["faiz"], s["bitis"], s["kalan_gun"], s["teminat"], s["teminat_tutar"], s["teminat_oran"], s["tl_limit"], s["tl_kullanim"]],
               para_s=(5, 6, 7, 8, 9, 13, 14, 15, 16), tarih_s=(10,), dolgu={7: KIRMIZI if s["bos"] < 0 else None, 8: KIRMIZI if s["oran"] > 100 else SARI if s["oran"] >= 90 else None,
                                                                             11: KIRMIZI if (s["kalan_gun"] or 0) < 0 else None})
    ws = wb.create_sheet("Banka Özeti")
    _baslik(ws, ["Banka", "Nakdi Limit", "Nakdi Risk", "Gayrinakdi Limit", "Gayrinakdi Risk", "Toplam Risk", "Boş Limit", "Pay %"], (22, 17, 17, 17, 17, 17, 17, 9))
    for b in r["bankalar"]:
        _satir(ws, [b["banka"], b["nakdi_limit"], b["nakdi"], b["gn_limit"], b["gn"], b["risk"], b["bos"], b["pay"]], para_s=range(2, 9))
    t = r["toplam"]
    n = _satir(ws, ["TOPLAM (TL)", t["nakdi_limit"], t["nakdi"], t["gn_limit"], t["gn"], t["risk"], t["bos"], None], para_s=range(2, 8))
    for h in ws[n]:
        h.font = Font(bold=True)
    ws.append([])
    for pb, f in r["faiz"].items():
        ws.append([f"Ağırlıklı ortalama faiz ({pb}, nakdi)", float(f)])
    ws.append([f"Durum tarihi {r['bugun']:%d.%m.%Y}. Tutarlar TL karşılığıdır; kuru verilmeyen para birimleri toplamlara girmez."])
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, kur_yol: Path | None = None, esik: Decimal = Decimal(90), gun: int = 30, yogunlasma: Decimal = Decimal(40)) -> dict:
    satirlar, uy = oku(yol)
    r = raporla(satirlar, kur_oku(kur_yol) if kur_yol else {"TL": Decimal(1)}, bugun, esik, gun, yogunlasma)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Banka bazında kredi limiti, kullanım, teminat ve faiz raporu üretir.")
    p.add_argument("--limitler", type=Path, help="Banka, Ürün, Sınıf, Para Birimi, Limit, Kullanım, Faiz Oranı, Limit Bitiş, Teminat, Teminat Tutarı")
    p.add_argument("--kurlar", type=Path, help="Para Birimi, Kur (dövizli satırlar için)")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--esik", default="90", help="Kullanım oranı uyarı eşiği, %% (varsayılan 90)")
    p.add_argument("--gun", type=int, default=30, help="Limit bitişine kaç gün kala uyarılsın (varsayılan 30)")
    p.add_argument("--yogunlasma", default="40", help="Banka payı uyarı eşiği, %% (varsayılan 40)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "banka_limit_risk.xlsx")
    a = p.parse_args(argv)
    ornek = a.limitler is None
    yol, kur_yol = (ORNEK / "limitler.csv", a.kurlar or ORNEK / "kurlar.csv") if ornek else (a.limitler, a.kurlar)
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    esik, yog = para(a.esik), para(a.yogunlasma)
    if bugun is None or esik is None or yog is None or not 0 < esik <= 100 or not 0 < yog <= 100 or not 0 <= a.gun <= 365:
        print("[X] --tarih GG.AA.YYYY biçiminde, --esik ve --yogunlasma 0–100 arasında, --gun 0–365 arasında olmalı")
        return 2
    for y in (yol, kur_yol):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(yol, a.cikti, bugun, kur_yol, esik, a.gun, yog)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {len(r['bankalar'])} banka · nakdi risk {t['nakdi']:,.2f} / limit {t['nakdi_limit']:,.2f} · gayrinakdi risk {t['gn']:,.2f} / limit {t['gn_limit']:,.2f} · boş limit {t['bos']:,.2f} TL")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
