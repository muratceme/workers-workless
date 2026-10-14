import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def bom(*satirlar):
    """(kod, ad, miktar, birim, fiyat[, üst])"""
    d = {}
    for s in satirlar:
        ust = s[5] if len(s) > 5 else "A"
        d[(main.katla(ust), main.katla(s[0]))] = {"ust": ust, "kod": s[0], "ad": s[1], "miktar": D(str(s[2])), "birim": s[3], "fiyat": None if s[4] is None else D(str(s[4]))}
    return d


class KarsilastirmaTesti(unittest.TestCase):
    def test_durumlar(self):
        eski = bom(("B1", "Vida", 4, "adet", 2), ("B2", "Somun", 4, "adet", 1), ("B3", "Sac", 1.5, "kg", 40), ("B4", "Boya", 0.2, "lt", 100), ("B5", "Etiket", 1, "adet", 0.5), ("B7", "Conta", 2, "adet", 3))
        yeni = bom(("b1", "Vida", 6, "ADET", 2), ("B3", "Sac", 1500, "g", 0.04), ("B4", "Boya", 0.2, "lt", 120), ("B5", "Etiket (yeni)", 1, "adet", 0.5), ("B6", "Kapak", 1, "adet", 10), ("B7", "Conta", 2, "adet", 3))
        r = main.karsilastir(eski, yeni)
        self.assertEqual([(s["kod"], s["durum"], s["eski_miktar"], s["yeni_miktar"], s["fark"], s["fark_oran"], s["eski_tutar"], s["yeni_tutar"], s["etki"]) for s in r["satirlar"]],
                         [("B6", "Eklendi", None, D(1), None, None, None, D(10), D(10)), ("B2", "Çıkarıldı", D(4), None, None, None, D(4), None, D(-4)), ("B3", "Birim değişti", D("1.5"), D(1500), None, None, D("60.0"), D("60.00"), D(0)),
                          ("b1", "Miktar değişti", D(4), D(6), D(2), D("50.00"), D(8), D(12), D(4)), ("B4", "Fiyat değişti", D("0.2"), D("0.2"), None, None, D("20.0"), D("24.0"), D("4.0")),
                          ("B5", "Tanım değişti", D(1), D(1), None, None, D("0.5"), D("0.5"), D(0)), ("B7", "Aynı", D(2), D(2), None, None, D(6), D(6), D(0))])
        self.assertEqual(r["ozet"], {"eski": 6, "yeni": 6, "Eklendi": 1, "Çıkarıldı": 1, "Birim değişti": 1, "Miktar değişti": 1, "Fiyat değişti": 1, "Tanım değişti": 1, "Aynı": 1, "eski_tutar": D("98.5"), "yeni_tutar": D("112.5"),
                                     "etki": D(14), "etki_oran": D("14.21"), "fiyatsiz": 0})
        self.assertEqual(([s["kod"] for s in r["degisenler"]], [(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]]), (["B6", "B2", "B3", "b1", "B4", "B5"], [("Orta", "Birim değişti", "A › B3")]))

    def test_fiyat_ve_ust_kod(self):
        r = main.karsilastir(bom(("X", "x", 2, "adet", 5), ("Y", "y", 1, "adet", None), ("Z", "z", 1, "adet", None), ("X", "x", 1, "adet", 5, "B")), bom(("X", "x", 3, "adet", None), ("Y", "y", 1, "adet", 7), ("Z", "z", 2, "adet", None), ("X", "x", 1, "adet", 5, "C")))
        self.assertEqual({(s["ust"], s["kod"]): (s["durum"], s["eski_fiyat"], s["yeni_fiyat"], s["etki"]) for s in r["satirlar"]},
                         {("A", "X"): ("Miktar değişti", D(5), D(5), D(5)), ("A", "Y"): ("Aynı", D(7), D(7), D(0)), ("A", "Z"): ("Miktar değişti", None, None, None), ("B", "X"): ("Çıkarıldı", D(5), None, D(-5)),
                          ("C", "X"): ("Eklendi", None, D(5), D(5))})
        self.assertEqual((r["ozet"]["etki"], r["ozet"]["fiyatsiz"], [u["tur"] for u in r["uyarilar"]]), (D(5), 1, ["Maliyeti olmayan kalemler"]))
        o = main.karsilastir(bom(("X", "x", 1, "adet", None)), bom(("X", "x", 1, "adet", None)))["ozet"]
        self.assertEqual((o["etki"], o["eski_tutar"], o["fiyatsiz"], o["Aynı"]), (None, None, 0, 1))
        self.assertEqual([main.miktar_oku(x) for x in ("0.250", "1,5", "1.250,5", "1.000", "12.5", 3, 0.5, "", "x", "0,125")], [D("0.250"), D("1.5"), D("1250.5"), D(1000), D("12.5"), D(3), D("0.5"), None, None, D("0.125")])
        with self.assertRaises(ValueError):
            main.karsilastir({}, {})


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "bom_v1.csv", ORNEK / "bom_v2.csv", cikti)
            self.assertEqual([(s["kod"], s["durum"], s["fark"], s["fark_oran"], s["etki"]) for s in r["degisenler"]],
                             [("HM-004", "Eklendi", None, None, D(4)), ("HM-022", "Çıkarıldı", None, None, D(-7)), ("HM-011", "Birim değişti", None, None, D(0)), ("HM-001", "Miktar değişti", D(-2), D("-25.00"), D(-5)),
                              ("HM-010", "Miktar değişti", D("-0.3"), D("-8.57"), D("-12.6")), ("HM-003", "Fiyat değişti", None, None, D(3)), ("HM-002", "Tanım değişti", None, None, D(0))])
            o = r["ozet"]
            self.assertEqual((o["eski"], o["yeni"], [o[d] for d in main.DURUM], o["eski_tutar"], o["yeni_tutar"], o["etki"], o["etki_oran"], o["fiyatsiz"]), (11, 11, [1, 1, 1, 2, 1, 1, 5], D("321.2"), D("303.6"), D("-17.6"), D("-5.48"), 2))
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]], [("Orta", "Birim değişti", "YM-110 › HM-011"), ("Bilgi", "Maliyeti olmayan kalemler", "Toplam")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Değişiklikler", "Tüm Kalemler", "Özet", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            e, y = Path(tmp) / "e.csv", Path(tmp) / "y.csv"
            e.write_text("Malzeme Kodu;Malzeme Adı;Miktar;Ölçü Birimi;Birim Fiyat\nK1;Vida;0.250;kg;1.250,50\nK2;Somun;2;adet;\nk1;Vida;0,5;kg;\nK3;X;0;adet;\nK4;X;y;adet;\nK5;X;1;adet;-2\n;X;1;adet;\n", encoding="utf-8")
            y.write_text("Kod;Miktar\nK1;1\n", encoding="utf-8")
            d, uy = main.oku(e, "Eski sürüm")
            self.assertEqual((d, [(u["tur"], u["kim"]) for u in uy]),
                             ({("", "k1"): {"ust": "", "kod": "K1", "ad": "Vida", "miktar": D("0.750"), "birim": "kg", "fiyat": D("1250.50")}, ("", "k2"): {"ust": "", "kod": "K2", "ad": "Somun", "miktar": D(2), "birim": "adet", "fiyat": None}},
                              [("Aynı kalem birden çok satırda", "k1"), ("Satır okunamadı", "K3"), ("Satır okunamadı", "K4"), ("Satır okunamadı", "K5")]))
            c = str(Path(tmp) / "x.xlsx")
            self.assertEqual(main.main(["--eski", str(e), "--yeni", str(y), "--cikti", c]), 0)
            self.assertEqual((main.main(["--eski", str(e)]), main.main(["--eski", str(e), "--yeni", str(Path(tmp) / "yok.csv")])), (2, 1))
            y.write_text("Bileşen Kodu;Bileşen Adı\nA;B\n", encoding="utf-8")
            self.assertEqual(main.main(["--eski", str(e), "--yeni", str(y), "--cikti", c]), 1)
            e.write_text("Bileşen Kodu;Miktar\n;1\n", encoding="utf-8")
            y.write_text("Bileşen Kodu;Miktar\n;1\n", encoding="utf-8")
            self.assertEqual(main.main(["--eski", str(e), "--yeni", str(y), "--cikti", c]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
