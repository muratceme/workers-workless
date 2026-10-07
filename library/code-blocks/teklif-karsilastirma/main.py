"""
Teklif Karşılaştırma — Workers / Workless kod bloğu
Satın Alma › Satın Alma Uzmanı

Tedarikçi tekliflerini tek tabloda toplar ve karşılaştırır:
  - Dövizli teklifleri verilen kurla TL'ye çevirir; nakliye ve ek masrafları toplam maliyete ekler.
  - Ödeme vadesi farkını, verilen yıllık iskonto (finansman) oranıyla bugünkü değere indirger.
  - Kalem bazında en uygun teklifi ve bölünmüş alım (her kalemi en uygun tedarikçiden) toplamını çıkarır.
  - Tedarikçileri fiyat, termin, vade, kalite ve garanti kriterleriyle ağırlıklı puanlar (en iyi = 100).
  - Teknik olarak uygun olmayan teklifleri ayırır; eksik kalemli teklifleri ve medyanın çok altındaki
    (aşırı düşük) fiyatları işaretler.
İnternete bağlanmaz.

Kullanım:
    python main.py                                          # örnek tekliflerle dener
    python main.py --girdi teklifler.xlsx --kur USD=41,20 EUR=48,05 --faiz 40
    python main.py --girdi teklifler.xlsx --agirlik fiyat=60 termin=10 vade=10 kalite=15 garanti=5
"""
from __future__ import annotations

import argparse
import csv
import sys
from collections import OrderedDict, defaultdict
from dataclasses import dataclass
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path
from statistics import median

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
KURUS = Decimal("0.01")
SIFIR = Decimal(0)
VARSAYILAN_AGIRLIK = {"fiyat": 50, "termin": 15, "vade": 15, "kalite": 15, "garanti": 5}


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def para(x) -> Decimal:
    if x in (None, ""):
        return SIFIR
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("₺", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return SIFIR


def yuvarla(x: Decimal) -> Decimal:
    return x.quantize(KURUS, rounding=ROUND_HALF_UP)


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ayirici = csv.Sniffer().sniff(metin.splitlines()[0], delimiters=",;\t").delimiter
        satirlar = list(csv.reader(metin.splitlines(), delimiter=ayirici))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


ALANLAR = {
    "tedarikci": ("tedarikçi", "firma", "teklif veren", "tedarikçi adı"),
    "kod": ("malzeme kodu", "stok kodu", "kalem no", "kod"),
    "malzeme": ("malzeme", "malzeme adı", "ürün", "kalem", "açıklama"),
    "miktar": ("miktar", "adet"),
    "birim": ("birim",),
    "fiyat": ("birim fiyat", "fiyat", "birim fiyatı"),
    "doviz": ("para birimi", "döviz", "pb", "döviz cinsi"),
    "nakliye": ("nakliye", "nakliye bedeli", "taşıma"),
    "ek": ("ek masraf", "diğer masraf", "ambalaj"),
    "termin": ("termin (gün)", "termin", "teslim süresi", "teslim süresi (gün)"),
    "vade": ("ödeme vadesi (gün)", "ödeme vadesi", "vade", "vade (gün)"),
    "garanti": ("garanti (ay)", "garanti", "garanti süresi"),
    "kalite": ("kalite puanı", "kalite", "tedarikçi puanı"),
    "uygun": ("teknik uygun", "teknik uygunluk", "uygun"),
    "not": ("not", "teklif notu"),
}


@dataclass
class Teklif:
    tedarikci: str
    kod: str
    malzeme: str
    miktar: Decimal
    birim: str
    fiyat: Decimal
    doviz: str
    nakliye: Decimal
    ek: Decimal
    termin: Decimal
    vade: Decimal
    garanti: Decimal
    kalite: Decimal | None
    uygun: bool
    not_: str
    tl_tutar: Decimal = SIFIR       # (birim fiyat × miktar) TL + nakliye + ek
    bd_tutar: Decimal = SIFIR       # vade iskontolu bugünkü değer

    @property
    def birim_tl(self) -> Decimal:
        return self.bd_tutar / self.miktar if self.miktar else SIFIR


def teklifleri_oku(yol: Path) -> list[Teklif]:
    s = tablo_oku(yol)
    b = [kucuk(x) for x in s[0]]
    k = {alan: next((i for i, x in enumerate(b) if x in adlar), None) for alan, adlar in ALANLAR.items()}
    if k["tedarikci"] is None or k["fiyat"] is None or (k["kod"] is None and k["malzeme"] is None):
        raise SystemExit(f"Tedarikçi, Malzeme (veya Malzeme Kodu) ve Birim Fiyat sütunları gerekli. Başlıklar: {s[0]}")
    al = lambda r, a: r[k[a]] if k[a] is not None and k[a] < len(r) else None  # noqa: E731
    teklifler = []
    for r in s[1:]:
        ted = str(al(r, "tedarikci") or "").strip()
        if not ted or al(r, "fiyat") in (None, ""):
            continue
        malzeme = str(al(r, "malzeme") or "").strip()
        uygun = kucuk(al(r, "uygun")) not in ("h", "hayır", "hayir", "no", "uygun değil", "0", "false")
        teklifler.append(Teklif(
            ted, str(al(r, "kod") or malzeme).strip(), malzeme, para(al(r, "miktar")) or Decimal(1), str(al(r, "birim") or "").strip(),
            para(al(r, "fiyat")), (str(al(r, "doviz") or "TL").strip().upper().replace("TRY", "TL") or "TL"),
            para(al(r, "nakliye")), para(al(r, "ek")), para(al(r, "termin")), para(al(r, "vade")), para(al(r, "garanti")),
            para(al(r, "kalite")) if al(r, "kalite") not in (None, "") else None, uygun, str(al(r, "not") or "").strip()))
    return teklifler


# ----------------------------------------------------------------------------
# Hesap
# ----------------------------------------------------------------------------

def bugunku_deger(tutar: Decimal, vade_gun: Decimal, yillik_faiz: Decimal) -> Decimal:
    """Basit iskonto: tutar ÷ (1 + yıllık oran × gün / 365). Vadesi uzun teklif bugünkü değerde ucuzlar."""
    return tutar / (1 + yillik_faiz * vade_gun / 365)


def az_iyi(deger: Decimal, en_iyi: Decimal) -> float:
    if deger <= 0:
        return 100.0
    return float(min(Decimal(100), en_iyi / deger * 100))


def cok_iyi(deger: Decimal, en_iyi: Decimal) -> float:
    return 100.0 if en_iyi <= 0 else float(deger / en_iyi * 100)


def calistir(girdi: Path, cikti: Path, kurlar: dict[str, Decimal] | None = None, yillik_faiz: Decimal = SIFIR,
             agirlik: dict | None = None, asiri_dusuk: Decimal = Decimal("0.30")) -> dict:
    kurlar = {"TL": Decimal(1)} | {k.upper(): v for k, v in (kurlar or {}).items()}
    agirlik = {**VARSAYILAN_AGIRLIK, **(agirlik or {})}
    toplam_agirlik = sum(agirlik.values())
    if toplam_agirlik <= 0:
        raise SystemExit("Kriter ağırlıklarının toplamı sıfırdan büyük olmalı.")
    teklifler = teklifleri_oku(girdi)
    eksik_kur = sorted({t.doviz for t in teklifler} - set(kurlar))
    if eksik_kur:
        raise SystemExit(f"Şu para birimleri için kur verin: {', '.join(eksik_kur)} (ör. --kur {eksik_kur[0]}=41,20)")
    uyarilar = []
    for t in teklifler:
        t.tl_tutar = t.fiyat * t.miktar * kurlar[t.doviz] + t.nakliye + t.ek
        t.bd_tutar = bugunku_deger(t.tl_tutar, t.vade, yillik_faiz)

    uygun = [t for t in teklifler if t.uygun]
    for t in teklifler:
        if not t.uygun:
            uyarilar.append(f"{t.tedarikci} · {t.malzeme or t.kod}: teknik olarak uygun değil, değerlendirme dışı")
    kalemler: "OrderedDict[str, list[Teklif]]" = OrderedDict()
    for t in teklifler:
        kalemler.setdefault(t.kod, [])
    for t in uygun:
        kalemler[t.kod].append(t)
    miktarlar = {}
    for kod in kalemler:
        m = {t.miktar for t in teklifler if t.kod == kod}
        miktarlar[kod] = max(m)
        if len(m) > 1:
            uyarilar.append(f"{kod}: tedarikçiler farklı miktarlara teklif vermiş ({', '.join(f'{x:g}' for x in sorted(m))}); "
                            "karşılaştırma birim fiyat üzerinden yapıldı")

    # Kalem bazında en iyi ve aşırı düşük kontrolü
    en_iyi: dict[str, Teklif] = {}
    for kod, liste in kalemler.items():
        if not liste:
            uyarilar.append(f"{kod}: uygun teklif yok")
            continue
        en_iyi[kod] = min(liste, key=lambda t: t.birim_tl)
        if len(liste) >= 3:
            med = Decimal(str(median([float(t.birim_tl) for t in liste])))
            for t in liste:
                if t.birim_tl < med * (1 - asiri_dusuk):
                    uyarilar.append(f"{t.tedarikci} · {t.malzeme or kod}: birim fiyat medyanın %{(1 - t.birim_tl / med) * 100:.0f} "
                                    "altında (aşırı düşük teklif?) — kapsam ve şartname uyumunu teyit edin")

    # Tedarikçi puanlaması
    tedarikciler: "OrderedDict[str, dict]" = OrderedDict()
    for t in uygun:
        d = tedarikciler.setdefault(t.tedarikci, {"teklifler": []})
        d["teklifler"].append(t)
    for ad, d in tedarikciler.items():
        ts = d["teklifler"]
        d["kapsam"] = len({t.kod for t in ts}) / len([k for k in kalemler if kalemler[k]])
        d["eksik"] = [k for k in kalemler if kalemler[k] and k not in {t.kod for t in ts}]
        d["tutar"] = sum((t.tl_tutar for t in ts), SIFIR)
        d["bd_tutar"] = sum((t.bd_tutar for t in ts), SIFIR)
        # Fiyat oranı: tedarikçinin teklif verdiği kalemlerde, en iyi birim fiyatlarla hesaplanan tutara oranı
        en_iyi_tutar = sum((en_iyi[t.kod].birim_tl * t.miktar for t in ts), SIFIR)
        d["fiyat_orani"] = d["bd_tutar"] / en_iyi_tutar if en_iyi_tutar else Decimal(1)
        d["termin"] = max(t.termin for t in ts)                  # sipariş, son kalem gelince tamamlanır
        d["vade"] = min(t.vade for t in ts)
        d["garanti"] = min(t.garanti for t in ts)
        kaliteler = [t.kalite for t in ts if t.kalite is not None]
        d["kalite"] = sum(kaliteler, SIFIR) / len(kaliteler) if kaliteler else None
    if not tedarikciler:
        raise SystemExit("Değerlendirilecek uygun teklif yok.")
    en_dusuk_termin = min(d["termin"] for d in tedarikciler.values())
    en_uzun_vade = max(d["vade"] for d in tedarikciler.values())
    en_uzun_garanti = max(d["garanti"] for d in tedarikciler.values())
    en_iyi_kalite = max((d["kalite"] for d in tedarikciler.values() if d["kalite"] is not None), default=None)
    for d in tedarikciler.values():
        p = {"fiyat": az_iyi(d["fiyat_orani"], Decimal(1)),
             "termin": az_iyi(d["termin"], en_dusuk_termin) if en_dusuk_termin > 0 else (100.0 if d["termin"] == 0 else 0.0),
             "vade": cok_iyi(d["vade"], en_uzun_vade),
             "kalite": 0.0 if d["kalite"] is None or en_iyi_kalite is None else cok_iyi(d["kalite"], en_iyi_kalite),
             "garanti": cok_iyi(d["garanti"], en_uzun_garanti)}
        d["puanlar"] = p
        d["toplam_puan"] = sum(p[k] * agirlik[k] for k in agirlik) / toplam_agirlik
    sira = sorted(tedarikciler.items(), key=lambda i: (-(i[1]["kapsam"] == 1), -i[1]["toplam_puan"]))
    tam = [(ad, d) for ad, d in sira if d["kapsam"] == 1]
    oneri_tek = tam[0][0] if tam else None
    bolunmus = sum((en_iyi[k].birim_tl * miktarlar[k] for k in en_iyi), SIFIR)
    sonuc = {"teklifler": teklifler, "kalemler": kalemler, "en_iyi": en_iyi, "tedarikciler": tedarikciler, "sira": sira,
             "oneri_tek": oneri_tek, "bolunmus": bolunmus, "miktarlar": miktarlar, "uyarilar": uyarilar,
             "agirlik": agirlik, "yillik_faiz": yillik_faiz, "kurlar": kurlar}
    _rapor(sonuc, cikti)
    return sonuc


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
YESIL = PatternFill("solid", fgColor="E3F5E1")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
GRI = PatternFill("solid", fgColor="EDEDED")
PARA = "#,##0.00"


def tl(x) -> str:
    return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def _rapor(s, cikti):
    wb = Workbook()
    o = wb.active
    o.title = "Değerlendirme"
    o.append(["Teklif Karşılaştırma ve Değerlendirme"])
    o["A1"].font = Font(bold=True, size=13)
    o.append(["Kriter ağırlıkları: " + " · ".join(f"{k} {v}" for k, v in s["agirlik"].items())
              + f"   |   vade iskontosu: yıllık %{s['yillik_faiz'] * 100:g}"
              + ("   |   kurlar: " + ", ".join(f"{k}={v}" for k, v in s["kurlar"].items() if k != "TL") if len(s["kurlar"]) > 1 else "")])
    o.append([])
    _baslik(o, ["Sıra", "Tedarikçi", "Kapsam", "Toplam (TL, KDV hariç)", "Bugünkü Değer (TL)", "Fiyat / En İyi", "Termin (gün)",
                "Vade (gün)", "Kalite", "Garanti (ay)", "Fiyat P.", "Termin P.", "Vade P.", "Kalite P.", "Garanti P.", "TOPLAM PUAN", "Not"],
            (6, 26, 9, 18, 18, 12, 11, 10, 9, 11, 9, 9, 9, 9, 9, 13, 40))
    for i, (ad, d) in enumerate(s["sira"], 1):
        p = d["puanlar"]
        o.append([i, ad, d["kapsam"], float(yuvarla(d["tutar"])), float(yuvarla(d["bd_tutar"])), float(d["fiyat_orani"]), float(d["termin"]),
                  float(d["vade"]), None if d["kalite"] is None else float(d["kalite"]), float(d["garanti"]),
                  round(p["fiyat"], 1), round(p["termin"], 1), round(p["vade"], 1), round(p["kalite"], 1), round(p["garanti"], 1),
                  round(d["toplam_puan"], 1), ("Eksik kalem: " + ", ".join(d["eksik"])) if d["eksik"] else ""])
        r = o.max_row
        o.cell(r, 3).number_format = "0%"
        o.cell(r, 4).number_format = o.cell(r, 5).number_format = PARA
        o.cell(r, 6).number_format = "0.000"
        o.cell(r, 16).font = Font(bold=True)
        if ad == s["oneri_tek"]:
            for c in range(1, 18):
                o.cell(r, c).fill = YESIL
        elif d["eksik"]:
            o.cell(r, 17).fill = GRI
    o.append([])
    tek = s["tedarikciler"].get(s["oneri_tek"]) if s["oneri_tek"] else None
    o.append(["Öneri", f"Tek tedarikçi: {s['oneri_tek']} (en yüksek puanlı, tüm kalemlere teklif veren)" if tek
              else "Tüm kalemlere teklif veren uygun tedarikçi yok; bölünmüş alım değerlendirin"])
    o.append(["", f"Bölünmüş alım (her kalem en uygun bugünkü değerden): {tl(s['bolunmus'])} TL"
              + (f" · tek tedarikçiye göre fark {tl(tek['bd_tutar'] - s['bolunmus'])} TL" if tek else "")])
    o.append(["", "Puanlar: fiyat, termin = en iyi / teklif × 100; vade, kalite, garanti = teklif / en iyi × 100"])
    for u in s["uyarilar"]:
        o.append(["Uyarı", u])
    o.freeze_panes = "C5"

    k = wb.create_sheet("Kalem Bazında")
    tedler = list(s["tedarikciler"])
    _baslik(k, ["Malzeme Kodu", "Malzeme", "Miktar", "Birim"] + [f"{t} (birim, TL)" for t in tedler] +
            ["En Uygun", "En Uygun Birim (TL)", "En Uygun Tutar (TL)", "İkinciye Göre Fark (TL)"],
            [14, 30, 9, 7] + [16] * len(tedler) + [22, 16, 17, 18])
    for kod, liste in s["kalemler"].items():
        if not liste:
            continue
        ornek = liste[0]
        fiyatlar = {t.tedarikci: t.birim_tl for t in liste}
        en = s["en_iyi"][kod]
        sirali = sorted(t.birim_tl for t in liste)
        fark = (sirali[1] - sirali[0]) * s["miktarlar"][kod] if len(sirali) > 1 else None
        k.append([kod, ornek.malzeme, float(s["miktarlar"][kod]), ornek.birim] +
                 [float(yuvarla(fiyatlar[t])) if t in fiyatlar else None for t in tedler] +
                 [en.tedarikci, float(yuvarla(en.birim_tl)), float(yuvarla(en.birim_tl * s["miktarlar"][kod])),
                  None if fark is None else float(yuvarla(fark))])
        r = k.max_row
        for j, t in enumerate(tedler, 5):
            k.cell(r, j).number_format = PARA
            if t == en.tedarikci:
                k.cell(r, j).fill = YESIL
            elif t not in fiyatlar:
                k.cell(r, j).fill = GRI
        for c in range(5 + len(tedler) + 1, 5 + len(tedler) + 4):
            k.cell(r, c).number_format = PARA
    k.freeze_panes = "C2"

    d = wb.create_sheet("Teklif Listesi")
    _baslik(d, ["Tedarikçi", "Malzeme Kodu", "Malzeme", "Miktar", "Birim Fiyat", "Para Birimi", "Nakliye + Ek", "Tutar (TL)",
                "Vade (gün)", "Bugünkü Değer (TL)", "Termin (gün)", "Garanti (ay)", "Kalite", "Teknik Uygun", "Not"],
            (24, 14, 30, 9, 13, 10, 13, 15, 10, 16, 11, 11, 8, 12, 30))
    for t in s["teklifler"]:
        d.append([t.tedarikci, t.kod, t.malzeme, float(t.miktar), float(t.fiyat), t.doviz, float(t.nakliye + t.ek), float(yuvarla(t.tl_tutar)),
                  float(t.vade), float(yuvarla(t.bd_tutar)), float(t.termin), float(t.garanti), None if t.kalite is None else float(t.kalite),
                  "Evet" if t.uygun else "HAYIR", t.not_])
        for c in (5, 7, 8, 10):
            d.cell(d.max_row, c).number_format = PARA
        if not t.uygun:
            d.cell(d.max_row, 14).fill = KIRMIZI
    d.freeze_panes = "B2"
    d.auto_filter.ref = d.dimensions

    b = wb.create_sheet("Bilgi")
    for satir in [["Toplam", "Birim fiyat × miktar × kur + nakliye + ek masraf (KDV hariç)"],
                  ["Bugünkü değer", "Toplam ÷ (1 + yıllık iskonto oranı × vade günü / 365): uzun vadeli teklif bugünkü değerde ucuzlar"],
                  ["Fiyat puanı", "Tedarikçinin teklif verdiği kalemlerde, en iyi birim fiyatlarla hesaplanan tutarın kendi tutarına oranı × 100"],
                  ["Termin / vade / garanti", "Tedarikçi bazında en uzun termin, en kısa vade ve en kısa garanti alınır (en kötü kalem)"],
                  ["Kapsam", "Teklif verilen kalem sayısı / uygun teklifi olan kalem sayısı; tek tedarikçi önerisi yalnız %100 kapsamdan seçilir"],
                  ["Aşırı düşük", "Kalemde en az 3 uygun teklif varsa, medyanın belirtilen oranından fazla altında kalan birim fiyat"],
                  ["Not", "Puanlama bir karar destek aracıdır; nihai seçim şartname uyumu, numune ve referanslarla birlikte yapılmalıdır"]]:
        b.append(satir)
    b.column_dimensions["A"].width = 22
    b.column_dimensions["B"].width = 120
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def anahtar_deger(liste: list[str], ad: str) -> dict[str, Decimal]:
    sonuc = {}
    for x in liste or []:
        if "=" not in x:
            raise SystemExit(f"{ad} 'ANAHTAR=DEĞER' biçiminde olmalı: {x}")
        k, v = x.split("=", 1)
        sonuc[k.strip().lower() if ad == "--agirlik" else k.strip().upper()] = para(v)
    return sonuc


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description="Tedarikçi tekliflerini fiyat, termin, vade, kalite ve garantiyle karşılaştırır.")
    ap.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "teklifler.csv",
                    help="Teklifler (.xlsx/.csv): Tedarikçi, Malzeme Kodu, Malzeme, Miktar, Birim Fiyat, Para Birimi, Termin, Vade, ...")
    ap.add_argument("--kur", nargs="*", default=[], metavar="PB=KUR", help="Döviz kurları, ör. USD=41,20 EUR=48,05")
    ap.add_argument("--faiz", type=Decimal, default=SIFIR, help="Vade farkı için yıllık iskonto oranı, %% (ör. 40)")
    ap.add_argument("--agirlik", nargs="*", default=[], metavar="KRİTER=AĞIRLIK",
                    help="fiyat, termin, vade, kalite, garanti ağırlıkları (varsayılan 50/15/15/15/5)")
    ap.add_argument("--asiri-dusuk", type=Decimal, default=Decimal(30), help="Aşırı düşük teklif eşiği, medyanın %% altı (varsayılan 30)")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "teklif_karsilastirma.xlsx")
    a = ap.parse_args(argv)
    kurlar = anahtar_deger(a.kur, "--kur")
    if a.girdi == BURASI / "ornek_veri" / "teklifler.csv":
        kurlar = kurlar or {"USD": Decimal("41.20"), "EUR": Decimal("48.05")}   # örnek kurlar
        a.faiz = a.faiz or Decimal(40)
    agirlik = {k: float(v) for k, v in anahtar_deger(a.agirlik, "--agirlik").items()}
    bilinmeyen = set(agirlik) - set(VARSAYILAN_AGIRLIK)
    if bilinmeyen:
        raise SystemExit(f"Bilinmeyen kriter: {', '.join(bilinmeyen)} (fiyat, termin, vade, kalite, garanti)")
    s = calistir(a.girdi, a.cikti, kurlar, a.faiz / 100, agirlik, a.asiri_dusuk / 100)
    for i, (ad, d) in enumerate(s["sira"], 1):
        print(f"[{i}] {ad:<26} puan {d['toplam_puan']:5.1f} · bugünkü değer {tl(d['bd_tutar']):>14} TL · kapsam %{d['kapsam'] * 100:.0f}")
    print(f"[OK] Öneri: {s['oneri_tek'] or 'tek tedarikçi yok'} · bölünmüş alım {tl(s['bolunmus'])} TL")
    for u in s["uyarilar"]:
        print(f"[!] {u}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
