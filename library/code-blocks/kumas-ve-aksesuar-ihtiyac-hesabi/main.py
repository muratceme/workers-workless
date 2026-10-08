"""
Kumaş ve Aksesuar İhtiyaç Hesabı — Workers / Workless kod bloğu
Tekstil ve Konfeksiyon › Planlama › Planlama Uzmanı

Sipariş adetlerinden (model × renk × beden) kumaş metrajını/kilosunu ve aksesuar adetlerini hesaplar:
  - Brüt ihtiyaç = sipariş adedi × (1 + fazla kesim %) × birim tüketim × (1 + fire %)
  - Reçete satırı renge bağlıysa (kumaş, fermuar, iplik) malzeme ürün rengiyle, bedene bağlıysa (beden etiketi,
    bedene göre tüketim) ürün bedeniyle ayrışır.
  - Malzeme kartı verilirse: metre ↔ kg dönüşümü (en × gramaj), top/rulo sayısı, ambalaj katına ve asgari
    sipariş miktarına yuvarlama; stok verilirse net ihtiyaç.
İnternete bağlanmaz.

Kullanım:
    python main.py                                                   # örnek verilerle dener
    python main.py --siparis siparis.xlsx --recete recete.xlsx
    python main.py --siparis siparis.xlsx --recete recete.xlsx --malzeme malzeme_karti.xlsx --stok stok.xlsx --fazla-kesim 3
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import OrderedDict, defaultdict
from decimal import ROUND_CEILING, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
SIFIR = Decimal(0)
BIR = Decimal(1)
YUZ = Decimal(100)


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def katla(s) -> str:
    s = kucuk(s)
    for a, b in zip("ıöüşğç", "iousgc"):
        s = s.replace(a, b)
    return re.sub(r"[^a-z0-9%²]+", " ", s).strip()


def sayi(x, varsayilan: Decimal | None = SIFIR) -> Decimal | None:
    if x in (None, ""):
        return varsayilan
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("%", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return varsayilan


def evet(x) -> bool:
    return katla(x) in ("evet", "e", "x", "1", "var", "true", "yes")


def yukari(x: Decimal, kat: Decimal) -> Decimal:
    """x'i kat'ın en yakın üst katına yuvarlar."""
    if not kat or kat <= 0:
        return x
    return (x / kat).to_integral_value(rounding=ROUND_CEILING) * kat


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


class Tablo:
    def __init__(self, yol: Path):
        s = tablo_oku(yol)
        self.ad, self.bas, self.satirlar = yol.name, s[0], s[1:]
        self.kb = [katla(x) for x in self.bas]

    def sutun(self, *adlar, zorunlu: str | None = None):
        i = next((self.kb.index(katla(a)) for a in adlar if katla(a) in self.kb), None)
        if i is None and zorunlu:
            raise SystemExit(f"{self.ad}: '{zorunlu}' sütunu bulunamadı. Başlıklar: {self.bas}")
        return i

    @staticmethod
    def al(r, i):
        v = r[i] if i is not None and i < len(r) else None
        return v.strip() if isinstance(v, str) else v


# ----------------------------------------------------------------------------
# Okuma
# ----------------------------------------------------------------------------

def siparis_oku(yol: Path) -> list[dict]:
    """Uzun biçim (Model, Renk, Beden, Adet) veya geniş biçim (Model, Renk, S, M, L, XL …)."""
    t = Tablo(yol)
    i_sip = t.sutun("sipariş no", "sipariş", "po", "po no")
    i_mod = t.sutun("model", "model kodu", "ürün", "artikel", "stil", zorunlu="Model")
    i_renk = t.sutun("renk", "renk kodu", "color")
    i_bed = t.sutun("beden", "size")
    i_adet = t.sutun("adet", "miktar", "sipariş adedi", "qty")
    satirlar = []
    if i_bed is not None and i_adet is not None:
        for r in t.satirlar:
            if t.al(r, i_mod) and sayi(t.al(r, i_adet)):
                satirlar.append({"siparis": str(t.al(r, i_sip) or ""), "model": str(t.al(r, i_mod)), "renk": str(t.al(r, i_renk) or "-"),
                                 "beden": str(t.al(r, i_bed)), "adet": sayi(t.al(r, i_adet))})
        return satirlar
    sabit = {i_sip, i_mod, i_renk} | {t.sutun("toplam", "total")}
    bedenler = [(i, str(b).strip()) for i, b in enumerate(t.bas) if i not in sabit and str(b or "").strip()]
    if not bedenler:
        raise SystemExit(f"{t.ad}: Beden ve Adet sütunları ya da beden başlıklı sütunlar (S, M, L …) gerekli.")
    for r in t.satirlar:
        if not t.al(r, i_mod):
            continue
        for i, b in bedenler:
            a = sayi(t.al(r, i))
            if a:
                satirlar.append({"siparis": str(t.al(r, i_sip) or ""), "model": str(t.al(r, i_mod)), "renk": str(t.al(r, i_renk) or "-"),
                                 "beden": b, "adet": a})
    return satirlar


def recete_oku(yol: Path) -> dict[str, list[dict]]:
    t = Tablo(yol)
    i = {k: t.sutun(*v) for k, v in {
        "model": ("model", "model kodu", "ürün", "artikel"), "kod": ("malzeme kodu", "malzeme", "kod"),
        "ad": ("malzeme adı", "açıklama", "ad"), "tur": ("tür", "malzeme türü", "tip"), "birim": ("birim", "ölçü birimi"),
        "tuketim": ("tüketim", "birim tüketim", "adet başına tüketim", "kullanım miktarı", "miktar"),
        "fire": ("fire %", "fire", "fire oranı"), "renk": ("renge bağlı", "renk bağımlı", "renk"),
        "beden": ("beden", "yalnız beden", "geçerli beden"), "bedene": ("bedene bağlı", "beden bağımlı"),
        "renk_eslem": ("malzeme rengi", "malzeme renk"),
    }.items()}
    for k, ad in (("model", "Model"), ("kod", "Malzeme Kodu"), ("tuketim", "Tüketim")):
        if i[k] is None:
            raise SystemExit(f"{t.ad}: '{ad}' sütunu bulunamadı. Başlıklar: {t.bas}")
    recete: dict[str, list[dict]] = defaultdict(list)
    for r in t.satirlar:
        if not t.al(r, i["model"]) or not t.al(r, i["kod"]):
            continue
        tur = katla(t.al(r, i["tur"]))
        recete[str(t.al(r, i["model"]))].append({
            "kod": str(t.al(r, i["kod"])), "ad": str(t.al(r, i["ad"]) or ""),
            "tur": "Kumaş" if tur.startswith(("kumas", "ana kumas", "astar", "fabric")) else "Aksesuar" if tur else "Aksesuar",
            "birim": str(t.al(r, i["birim"]) or ""), "tuketim": sayi(t.al(r, i["tuketim"])), "fire": sayi(t.al(r, i["fire"])),
            "renge_bagli": evet(t.al(r, i["renk"])), "beden": str(t.al(r, i["beden"]) or "").strip(),
            "bedene_bagli": evet(t.al(r, i["bedene"])), "malzeme_rengi": str(t.al(r, i["renk_eslem"]) or "").strip()})
    return dict(recete)


def malzeme_oku(yol: Path | None) -> dict[str, dict]:
    if not yol:
        return {}
    t = Tablo(yol)
    i = {k: t.sutun(*v) for k, v in {
        "kod": ("malzeme kodu", "kod"), "ad": ("malzeme adı", "ad", "açıklama"),
        "ambalaj": ("ambalaj miktarı", "paket", "paket içi", "ambalaj"), "min": ("asgari sipariş", "min sipariş", "moq"),
        "top": ("top boyu", "top", "rulo", "top miktarı"), "en": ("en cm", "en", "kumaş eni"),
        "gramaj": ("gramaj g m²", "gramaj g m2", "gramaj"), "tedarikci": ("tedarikçi", "firma"),
        "termin": ("tedarik süresi gün", "tedarik süresi", "termin gün"),
    }.items()}
    if i["kod"] is None:
        raise SystemExit(f"{t.ad}: 'Malzeme Kodu' sütunu bulunamadı. Başlıklar: {t.bas}")
    return {str(t.al(r, i["kod"])): {"ad": str(t.al(r, i["ad"]) or ""), "ambalaj": sayi(t.al(r, i["ambalaj"]), None),
                                      "min": sayi(t.al(r, i["min"]), None), "top": sayi(t.al(r, i["top"]), None),
                                      "en": sayi(t.al(r, i["en"]), None), "gramaj": sayi(t.al(r, i["gramaj"]), None),
                                      "tedarikci": str(t.al(r, i["tedarikci"]) or ""), "termin": sayi(t.al(r, i["termin"]), None)}
            for r in t.satirlar if t.al(r, i["kod"])}


def stok_oku(yol: Path | None) -> dict[tuple[str, str, str], Decimal]:
    if not yol:
        return {}
    t = Tablo(yol)
    i_kod = t.sutun("malzeme kodu", "kod", zorunlu="Malzeme Kodu")
    i_renk, i_bed = t.sutun("renk", "malzeme rengi"), t.sutun("beden")
    i_m = t.sutun("stok", "mevcut stok", "miktar", "kullanılabilir stok", zorunlu="Stok")
    stok: dict[tuple[str, str, str], Decimal] = defaultdict(Decimal)
    for r in t.satirlar:
        if t.al(r, i_kod):
            stok[(str(t.al(r, i_kod)), str(t.al(r, i_renk) or "-"), str(t.al(r, i_bed) or "-"))] += sayi(t.al(r, i_m))
    return dict(stok)


# ----------------------------------------------------------------------------
# Hesap
# ----------------------------------------------------------------------------

def kg_metre(birim: str, miktar: Decimal, kart: dict | None) -> tuple[Decimal | None, Decimal | None]:
    """(metre, kg). Kartta en (cm) ve gramaj (g/m²) varsa diğer birim hesaplanır: kg/m = en/100 × gramaj/1000."""
    b = katla(birim)
    kgm = (kart["en"] / YUZ * kart["gramaj"] / 1000) if kart and kart.get("en") and kart.get("gramaj") else None
    if b in ("m", "mt", "metre"):
        return miktar, (miktar * kgm if kgm else None)
    if b in ("kg", "kilo", "kilogram"):
        return (miktar / kgm if kgm else None), miktar
    return None, None


def calistir(siparis_yolu: Path, recete_yolu: Path, cikti: Path, malzeme_yolu: Path | None = None,
             stok_yolu: Path | None = None, fazla_kesim: float = 0.0) -> dict:
    siparis = siparis_oku(siparis_yolu)
    recete = recete_oku(recete_yolu)
    kartlar, stok = malzeme_oku(malzeme_yolu), stok_oku(stok_yolu)
    fazla = Decimal(str(fazla_kesim))
    uyarilar = []
    recetesiz = sorted({s["model"] for s in siparis if s["model"] not in recete})
    if recetesiz:
        uyarilar.append(f"Reçetesi olmayan modeller hesaba katılmadı: {', '.join(recetesiz)}")
    ihtiyac: dict[tuple, dict] = OrderedDict()
    detay = []
    for s in siparis:
        uretim = s["adet"] * (BIR + fazla / YUZ)
        for r in recete.get(s["model"], []):
            if r["beden"] and katla(r["beden"]) != katla(s["beden"]):
                continue
            renk = (r["malzeme_rengi"] or s["renk"]) if r["renge_bagli"] else "-"
            beden = s["beden"] if r["bedene_bagli"] else "-"
            miktar = uretim * r["tuketim"] * (BIR + r["fire"] / YUZ)
            anahtar = (r["kod"], renk, beden)
            x = ihtiyac.setdefault(anahtar, {"kod": r["kod"], "ad": kartlar.get(r["kod"], {}).get("ad") or r["ad"],
                                            "tur": r["tur"], "birim": r["birim"], "renk": renk, "beden": beden,
                                            "brut": SIFIR, "net_tuketim": SIFIR, "modeller": set()})
            if x["birim"] and r["birim"] and katla(x["birim"]) != katla(r["birim"]):
                uyarilar.append(f"{r['kod']}: reçetelerde farklı birimler ({x['birim']} / {r['birim']}) — kontrol edin")
            x["brut"] += miktar
            x["net_tuketim"] += s["adet"] * r["tuketim"]
            x["modeller"].add(s["model"])
            detay.append({**s, "uretim": uretim, "kod": r["kod"], "renk_m": renk, "beden_m": beden, "tuketim": r["tuketim"],
                          "fire": r["fire"], "miktar": miktar, "birim": r["birim"]})
    for (kod, renk, beden), x in ihtiyac.items():
        kart = kartlar.get(kod)
        x["stok"] = stok.get((kod, renk, beden), SIFIR)
        x["net"] = max(SIFIR, x["brut"] - x["stok"])
        siparis_mik = x["net"]
        if x["tur"] == "Aksesuar" or katla(x["birim"]) in ("adet", "ad", "pcs"):
            siparis_mik = siparis_mik.to_integral_value(rounding=ROUND_CEILING)
        if kart:
            if kart.get("ambalaj"):
                siparis_mik = yukari(siparis_mik, kart["ambalaj"])
            if kart.get("min") and SIFIR < siparis_mik < kart["min"]:
                siparis_mik = kart["min"]
        x["siparis"] = siparis_mik
        x["metre"], x["kg"] = kg_metre(x["birim"], x["brut"], kart)
        x["top"] = (math.ceil(x["net"] / kart["top"]) if kart and kart.get("top") and x["net"] else None)
        x["tedarikci"] = (kart or {}).get("tedarikci", "")
        x["termin"] = (kart or {}).get("termin")
        if kod not in kartlar and kartlar:
            uyarilar.append(f"{kod}: malzeme kartında yok (ambalaj/top yuvarlaması yapılmadı)")
    toplam_adet = sum((s["adet"] for s in siparis), SIFIR)
    _rapor(siparis, ihtiyac, detay, uyarilar, fazla, cikti)
    return {"siparis": siparis, "ihtiyac": ihtiyac, "detay": detay, "uyarilar": list(dict.fromkeys(uyarilar)),
            "toplam_adet": toplam_adet}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
KUMAS_DOLGU = PatternFill("solid", fgColor="E8F0FE")
MIK = "#,##0.00"


def _baslik(ws, satir=1):
    for c in ws[satir]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI


def _f(x):
    return None if x is None else float(x)


def _rapor(siparis, ihtiyac, detay, uyarilar, fazla, cikti):
    wb = Workbook()
    o = wb.active
    o.title = "Özet"
    o.append(["Kumaş ve aksesuar ihtiyaç hesabı"])
    o["A1"].font = Font(bold=True, size=12)
    o.append(["Sipariş satırı", len(siparis)])
    o.append(["Toplam sipariş adedi", float(sum((s["adet"] for s in siparis), SIFIR))])
    o.append(["Fazla kesim payı", f"%{fazla:g}"])
    o.append(["Malzeme kalemi (kod × renk × beden)", len(ihtiyac)])
    o.append([])
    o.append(["Model", "Renk", "Adet"])
    _baslik(o, o.max_row)
    mr = defaultdict(Decimal)
    for s in siparis:
        mr[(s["model"], s["renk"])] += s["adet"]
    for (m, r), a in sorted(mr.items()):
        o.append([m, r, float(a)])
    for u in uyarilar:
        o.append(["Uyarı", u])
    o.column_dimensions["A"].width = 36
    o.column_dimensions["B"].width = 18

    k = wb.create_sheet("İhtiyaç Listesi")
    k.append(["Tür", "Malzeme Kodu", "Malzeme Adı", "Renk", "Beden", "Birim", "Net Tüketim (firesiz)", "Brüt İhtiyaç",
              "Metre", "Kg", "Stok", "Net İhtiyaç", "Top/Rulo", "Sipariş Miktarı (yuvarlanmış)", "Tedarikçi", "Tedarik Süresi (gün)",
              "Kullanıldığı Modeller"])
    _baslik(k)
    for x in sorted(ihtiyac.values(), key=lambda x: (x["tur"] != "Kumaş", x["kod"], x["renk"], x["beden"])):
        k.append([x["tur"], x["kod"], x["ad"], x["renk"], x["beden"], x["birim"], _f(x["net_tuketim"]), _f(x["brut"]),
                  _f(x["metre"]), _f(x["kg"]), _f(x["stok"]) or None, _f(x["net"]), x["top"], _f(x["siparis"]), x["tedarikci"],
                  _f(x["termin"]), ", ".join(sorted(x["modeller"]))])
        for c in (7, 8, 9, 10, 11, 12, 14):
            k.cell(k.max_row, c).number_format = MIK
        if x["tur"] == "Kumaş":
            for c in range(1, 18):
                k.cell(k.max_row, c).fill = KUMAS_DOLGU
    for j, w in enumerate((9, 14, 28, 12, 7, 7, 13, 13, 11, 11, 10, 12, 9, 15, 18, 10, 24), 1):
        k.column_dimensions[get_column_letter(j)].width = w
    k.freeze_panes = "C2"
    k.auto_filter.ref = k.dimensions

    d = wb.create_sheet("Hesap Detayı")
    d.append(["Sipariş No", "Model", "Renk", "Beden", "Sipariş Adedi", "Üretim Adedi (fazla kesim dahil)", "Malzeme Kodu",
              "Malzeme Rengi", "Malzeme Bedeni", "Birim Tüketim", "Fire %", "İhtiyaç", "Birim"])
    _baslik(d)
    for x in detay:
        d.append([x["siparis"], x["model"], x["renk"], x["beden"], _f(x["adet"]), _f(x["uretim"]), x["kod"], x["renk_m"], x["beden_m"],
                  _f(x["tuketim"]), _f(x["fire"]), _f(x["miktar"]), x["birim"]])
        d.cell(d.max_row, 12).number_format = MIK
    for j, w in enumerate((10, 12, 12, 7, 10, 14, 14, 12, 9, 10, 7, 12, 7), 1):
        d.column_dimensions[get_column_letter(j)].width = w
    d.freeze_panes = "B2"
    d.auto_filter.ref = d.dimensions

    b = wb.create_sheet("Bilgi")
    for s in [["Brüt ihtiyaç", "sipariş adedi × (1 + fazla kesim %) × birim tüketim × (1 + fire %)"],
              ["Net ihtiyaç", "brüt ihtiyaç − mevcut stok (aynı malzeme, renk ve beden)"],
              ["Sipariş miktarı", "aksesuarlarda tam sayıya, malzeme kartındaki ambalaj miktarının katına yuvarlanır; asgari sipariş miktarının altındaysa asgariye çıkarılır"],
              ["Metre ↔ kg", "kartta en (cm) ve gramaj (g/m²) varsa: kg/metre = en/100 × gramaj/1000"],
              ["Top/Rulo", "net ihtiyaç ÷ top boyu (yukarı yuvarlanır)"],
              ["Not", "Birim tüketimi pastal verimi ve kalıp ölçüsüne göre güncel tutun; boyahane/yıkama çekmesi ve kumaş hatası "
                      "paylarını fire oranına dahil edin."]]:
        b.append(s)
    b.column_dimensions["A"].width = 18
    b.column_dimensions["B"].width = 130
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ornek = BURASI / "ornek_veri"
    ap = argparse.ArgumentParser(description="Sipariş adetlerinden kumaş metrajını/kilosunu ve aksesuar adetlerini hesaplar.")
    ap.add_argument("--siparis", type=Path, default=ornek / "siparis.csv",
                    help="Sipariş (.xlsx/.csv): Model, Renk, Beden, Adet — veya Model, Renk, S, M, L … (geniş biçim)")
    ap.add_argument("--recete", type=Path, default=ornek / "recete.csv",
                    help="Reçete: Model, Malzeme Kodu, Malzeme Adı, Tür, Birim, Tüketim [, Fire %%, Renge Bağlı, Bedene Bağlı, Beden, Malzeme Rengi]")
    ap.add_argument("--malzeme", type=Path, help="Malzeme kartı: Malzeme Kodu [, Ambalaj Miktarı, Asgari Sipariş, Top Boyu, En (cm), Gramaj, Tedarikçi]")
    ap.add_argument("--stok", type=Path, help="Stok: Malzeme Kodu, [Renk, Beden,] Stok")
    ap.add_argument("--fazla-kesim", type=float, default=0.0, help="Sipariş adedine eklenecek fazla kesim/üretim payı, %% (varsayılan 0)")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "kumas_aksesuar_ihtiyaci.xlsx")
    a = ap.parse_args(argv)
    if a.siparis == ornek / "siparis.csv" and a.recete == ornek / "recete.csv":
        a.malzeme = a.malzeme or ornek / "malzeme_karti.csv"
        a.stok = a.stok or ornek / "stok.csv"
        if a.fazla_kesim == 0.0:
            a.fazla_kesim = 3.0
    s = calistir(a.siparis, a.recete, a.cikti, a.malzeme, a.stok, a.fazla_kesim)
    kumas = [x for x in s["ihtiyac"].values() if x["tur"] == "Kumaş"]
    print(f"[OK] {len(s['siparis'])} sipariş satırı · {s['toplam_adet']:g} adet · {len(s['ihtiyac'])} malzeme kalemi "
          f"({len(kumas)} kumaş)")
    for x in kumas:
        ek = f" · {x['kg']:.1f} kg" if x["kg"] is not None and katla(x["birim"]) != "kg" else ""
        print(f"    {x['kod']} {x['renk']}: brüt {x['brut']:.1f} {x['birim']}{ek} · net {x['net']:.1f} · sipariş {x['siparis']:.1f}")
    for u in s["uyarilar"]:
        print(f"[!] {u}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
