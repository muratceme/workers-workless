import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
DENETIM, BUGUN = date(2026, 11, 20), date(2026, 10, 14)


def m(std, madde, durum, kritik=False, guncelleme=None, termin=None, sorumlu="A", bolum="B1"):
    return {"standart": std, "bolum": bolum, "madde": madde, "gereklilik": "", "kayit": "", "kritik": kritik, "sorumlu": sorumlu, "durum": durum, "girilen": durum,
            "guncelleme": guncelleme, "termin": termin, "not": ""}


class AnalizTesti(unittest.TestCase):
    def test_puan(self):
        l = [m("X", "1", "Hazır", guncelleme=date(2026, 1, 1)), m("X", "2", "Hazır", True, guncelleme=date(2025, 11, 1)), m("X", "3", "Kısmen", termin=date(2026, 11, 1)),
             m("X", "4", "Eksik", True, termin=date(2026, 12, 1)), m("X", "5", "Uygulanamaz"), m("Y", "1", "Eksik", termin=date(2026, 10, 1)), m("Y", "2", "Kısmen", sorumlu="", bolum="B2")]
        r = main.analiz(l, DENETIM, BUGUN)
        g = r["genel"]
        # puan: 1 + 0,5 (eski doküman) + 0,5 + 0 + 0 + 0,5 = 2,5 / 6
        self.assertEqual((g["uyg"], g["Hazır"], g["Güncellenmeli"], g["oran"], g["kritik_eksik"], r["kalan"]), (6, 1, 1, D("41.67"), 2, 37))
        self.assertEqual([(s["ad"], s["oran"]) for s in r["standart"]], [("X", D("50.00")), ("Y", D("25.00"))])
        self.assertEqual([s["ad"] for s in r["bolum"]], [("Y", "B1"), ("X", "B1"), ("Y", "B2")])
        self.assertEqual([(a["standart"], a["madde"], a["kalan"]) for a in r["aksiyon"]], [("X", "4", -11), ("X", "2", None), ("Y", "1", 50), ("X", "3", 19), ("Y", "2", None)])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Orta", "Aksiyon sahipsiz", "X 2"), ("Orta", "Aksiyon sahipsiz", "Y 2"), ("Orta", "Doküman / kayıt eski", "X 2"), ("Orta", "Termin geçmiş", "Y 1"),
                          ("Yüksek", "Kritik madde hazır değil", "X 2"), ("Yüksek", "Kritik madde hazır değil", "X 4"), ("Yüksek", "Termin denetimden sonra", "X 4")])
        self.assertEqual(main.analiz(l, DENETIM, BUGUN, D(0), 1000)["genel"]["oran"], D("33.33"))
        self.assertEqual(main.analiz(l, date(2026, 1, 1), BUGUN)["uyarilar"][0]["tur"], "Denetim tarihi geçmiş")
        with self.assertRaises(ValueError):
            main.analiz([], DENETIM, BUGUN)

    def test_durum(self):
        self.assertEqual([main._durum(x) for x in ("Tamam", "taslak", "Revize", "", "Kapsam dışı", "x")], ["Hazır", "Kısmen", "Güncellenmeli", "Eksik", "Uygulanamaz", None])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "denetim_listesi.csv", cikti, DENETIM, BUGUN)
            g = r["genel"]
            self.assertEqual((g["madde"], g["uyg"], g["oran"], g["kritik_eksik"], len(r["aksiyon"])), (29, 28, D("62.50"), 9, 15))
            self.assertEqual([(s["ad"], s["oran"]) for s in r["standart"]], [("Müşteri denetimi", D("50.00")), ("Resmî kontrol", D("70.00")), ("Sertifikasyon", D("62.50"))])
            self.assertEqual(sorted(u["kim"] for u in r["uyarilar"] if u["tur"] == "Termin denetimden sonra"), ["Müşteri denetimi M-3", "Sertifikasyon S-2.4"])
            self.assertEqual([u["kim"] for u in r["uyarilar"] if u["tur"] == "Doküman / kayıt eski"], ["Sertifikasyon S-2.2"])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Aksiyon Listesi", "Tüm Maddeler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "l.csv"
            y.write_text("Denetim;Madde;Hazırlık Durumu;Termin;Kritik\nA;1;Var;;Evet\nA;1;Yok;;\nA;2;belki;;\nA;3;Yok;x;\nB;1;;;\n;;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["standart"], x["madde"], x["durum"], x["kritik"]) for x in l], [u["tur"] for u in uy]),
                             ([("A", "1", "Hazır", True), ("B", "1", "Eksik", False)], ["Madde iki kez yazılmış", "Satır okunamadı", "Satır okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--liste", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--denetim", "01.12.2026", "--bugun", "01.11.2026", "--kismen", "40", "--gecerlilik", "180"]), 0)
            self.assertEqual([main.main(arg + v) for v in ([], ["--denetim", "x"], ["--denetim", "01.12.2026", "--kismen", "101"], ["--denetim", "01.12.2026", "--gecerlilik", "0"])], [2] * 4)
            self.assertEqual(main.main(["--liste", str(Path(tmp) / "yok.csv"), "--denetim", "01.12.2026"]), 1)
            y.write_text("Madde;Durum\n1;belki\n", encoding="utf-8")
            self.assertEqual(main.main(arg + ["--denetim", "01.12.2026"]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
