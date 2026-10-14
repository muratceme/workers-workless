"""
Denetim Hazırlık Kontrol Listesi — Workers / Workless kod bloğu
Kalite Güvence ve Gıda Güvenliği › Kalite Güvence Uzmanı
BRCGS, IFS, resmî kontrol ya da müşteri denetimi maddelerine göre doküman ve kayıt hazırlık durumunu takip eder.
  Puan          = Hazır 1 · Kısmen / Güncellenmeli --kismen (%50) · Eksik 0 · Uygulanamaz hesaba girmez
  Hazırlık %    = puan toplamı ÷ uygulanabilir madde × 100 (standart, bölüm ve genel)
  Eski doküman  = "Hazır" maddenin son güncelleme tarihi denetimden --gecerlilik (365) günden eski → Güncellenmeli sayılır
  Aksiyon       = hazır olmayan maddeler; termin, sorumlu ve denetime kalan gün
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
SUTUN = {"standart": ("standart", "denetim", "denetim turu", "sema"), "bolum": ("bolum", "baslik", "alan", "kategori"), "madde": ("madde no", "madde", "no", "kod"),
         "gereklilik": ("gereklilik", "gereklilik ozeti", "aciklama", "konu"), "kayit": ("gerekli kayit", "gerekli dokuman", "kayit dokuman", "dokuman", "kanit"),
         "kritik": ("kritik", "temel", "fundamental", "ko"), "sorumlu": ("sorumlu", "sorumlu kisi"), "durum": ("durum", "hazirlik durumu"),
         "guncelleme": ("son guncelleme", "guncelleme tarihi", "revizyon tarihi", "kayit tarihi"), "termin": ("termin", "hedef tarih", "son tarih"), "not": ("not", "notlar", "yorum")}
DURUMLAR = ("Hazır", "Kısmen", "Güncellenmeli", "Eksik", "Uygulanamaz")


def _durum(x) -> str | None:
    k = katla(x)
    if k in ("hazir", "tamam", "uygun", "var", "evet", "karsilaniyor"):
        return "Hazır"
    if k in ("kismen", "kismi", "devam", "devam ediyor", "taslak"):
        return "Kısmen"
    if k in ("guncellenmeli", "revize", "revizyon", "guncelleme gerekli", "eski"):
        return "Güncellenmeli"
    if not k or k in ("eksik", "yok", "hayir", "hazir degil", "baslanmadi"):
        return "Eksik"
    if k in ("uygulanamaz", "dd", "kapsam disi", "na", "n a", "gecerli degil"):
        return "Uygulanamaz"
    return None


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("madde", "durum")):
        madde = metin(r["madde"])
        if not madde:
            continue
        d = _durum(r["durum"])
        g, t = (tarih(r[a]) if metin(r[a]) else None for a in ("guncelleme", "termin"))
        std = metin(r["standart"]) or "Genel"
        kim = f"{std} {madde}"
        if d is None or (metin(r["guncelleme"]) and g is None) or (metin(r["termin"]) and t is None):
            uy.append(uyari("Yüksek", "Satır okunamadı", kim, f"Satır {r['_satir']}: durum ('{metin(r['durum'])}'), son güncelleme ya da termin okunamadı."))
            continue
        if (katla(std), katla(madde)) in gorulen:
            uy.append(uyari("Orta", "Madde iki kez yazılmış", kim, f"Satır {r['_satir']}: ikinci satır alınmadı."))
            continue
        gorulen.add((katla(std), katla(madde)))
        l.append({"standart": std, "bolum": metin(r["bolum"]) or "Genel", "madde": madde, "gereklilik": metin(r["gereklilik"]), "kayit": metin(r["kayit"]),
                  "kritik": katla(r["kritik"]) in ("evet", "e", "x", "kritik", "temel", "fundamental", "ko", "1"), "sorumlu": metin(r["sorumlu"]), "durum": d, "girilen": d,
                  "guncelleme": g, "termin": t, "not": metin(r["not"])})
    return l, uy


def analiz(maddeler: list[dict], denetim: date, bugun: date, kismen: Decimal = Decimal(50), gecerlilik: int = 365) -> dict:
    if not maddeler:
        raise ValueError("Okunabilen madde yok.")
    uy = []
    kalan = (denetim - bugun).days
    if kalan < 0:
        uy.append(uyari("Bilgi", "Denetim tarihi geçmiş", f"{denetim:%d.%m.%Y}", "Denetim tarihi bugünden önce; liste denetim sonrası durum olarak okunmalı."))
    puan = {"Hazır": Decimal(1), "Kısmen": kismen / 100, "Güncellenmeli": kismen / 100, "Eksik": Decimal(0)}
    for m in maddeler:
        kim = f"{m['standart']} {m['madde']}"
        m["durum"] = m["girilen"]
        if m["durum"] == "Hazır" and m["guncelleme"] and (denetim - m["guncelleme"]).days > gecerlilik:
            m["durum"] = "Güncellenmeli"
            uy.append(uyari("Orta", "Doküman / kayıt eski", kim, f"Son güncelleme {m['guncelleme']:%d.%m.%Y}; denetim tarihinde {gecerlilik} günden eski olacak. Gözden geçirip tarihini yenileyin."))
        m["puan"] = puan.get(m["durum"])
        if m["durum"] in ("Hazır", "Uygulanamaz"):
            continue
        if m["kritik"]:
            uy.append(uyari("Yüksek", "Kritik madde hazır değil", kim, f"{m['gereklilik'] or m['kayit']} · durum {m['durum']}" + (f" · sorumlu {m['sorumlu']}" if m["sorumlu"] else "")))
        if m["termin"] and m["termin"] > denetim:
            uy.append(uyari("Yüksek", "Termin denetimden sonra", kim, f"Termin {m['termin']:%d.%m.%Y}, denetim {denetim:%d.%m.%Y}; denetime yetişmiyor."))
        elif m["termin"] and m["termin"] < bugun:
            uy.append(uyari("Orta", "Termin geçmiş", kim, f"Termin {m['termin']:%d.%m.%Y} geçti; durum {m['durum']}."))
        if not m["termin"] or not m["sorumlu"]:
            eksik = " ve ".join(a for a, v in (("termin", m["termin"]), ("sorumlu", m["sorumlu"])) if not v)
            uy.append(uyari("Orta", "Aksiyon sahipsiz", kim, f"Hazır olmayan maddede {eksik} yazılmamış."))

    def ozet(fn):
        g = {}
        for m in maddeler:
            g.setdefault(fn(m), []).append(m)
        satir = []
        for ad, l in g.items():
            uyg = [m for m in l if m["durum"] != "Uygulanamaz"]
            s = {"ad": ad, "madde": len(l), "uyg": len(uyg), **{d: sum(1 for m in l if m["durum"] == d) for d in DURUMLAR},
                 "oran": yuvarla(sum((m["puan"] for m in uyg), Decimal(0)) * 100 / len(uyg)) if uyg else None,
                 "kritik_eksik": sum(1 for m in uyg if m["kritik"] and m["durum"] != "Hazır")}
            satir.append(s)
        return satir
    standart = sorted(ozet(lambda m: m["standart"]), key=lambda s: katla(s["ad"]))
    bolum = sorted(ozet(lambda m: (m["standart"], m["bolum"])), key=lambda s: (s["oran"] is None, s["oran"] or 0, katla(s["ad"][0]), katla(s["ad"][1])))
    genel = ozet(lambda m: "GENEL")[0]
    aksiyon = sorted((m for m in maddeler if m["durum"] not in ("Hazır", "Uygulanamaz")),
                     key=lambda m: (not m["kritik"], m["termin"] is None, m["termin"] or date.max, katla(m["standart"]), katla(m["madde"])))
    for m in aksiyon:
        m["kalan"] = (denetim - m["termin"]).days if m["termin"] else None
    return {"standart": standart, "bolum": bolum, "genel": genel, "aksiyon": aksiyon, "maddeler": maddeler, "kalan": kalan, "denetim": denetim, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    _baslik(ws, ["Standart / Bölüm", "Bölüm", "Madde", "Uygulanabilir", "Hazır", "Kısmen", "Güncellenmeli", "Eksik", "Uygulanamaz", "Hazırlık %", "Kritik Hazır Değil"],
            (20, 26, 7, 11, 7, 7, 11, 7, 11, 10, 11))

    def yaz(a, b, s, kalin=False):
        n = _satir(ws, [a, b, s["madde"], s["uyg"], *(s[d] for d in DURUMLAR), s["oran"], s["kritik_eksik"]], para_s=(10,),
                   dolgu={10: YESIL if s["oran"] is not None and s["oran"] >= 90 else (SARI if s["oran"] is not None and s["oran"] >= 70 else KIRMIZI), 11: KIRMIZI if s["kritik_eksik"] else None})
        if kalin:
            for h in ws[n]:
                h.font = Font(bold=True)
    yaz("GENEL", f"Denetim {r['denetim']:%d.%m.%Y} · {r['kalan']} gün kaldı", r["genel"], True)
    for s in r["standart"]:
        yaz(s["ad"], "", s, True)
    for s in r["bolum"]:
        yaz(s["ad"][0], s["ad"][1], s)
    ws = wb.create_sheet("Aksiyon Listesi")
    _baslik(ws, ["Standart", "Madde", "Bölüm", "Gereklilik", "Gerekli Kayıt / Doküman", "Kritik", "Durum", "Sorumlu", "Termin", "Termin → Denetim (gün)", "Not"],
            (14, 9, 20, 36, 32, 7, 13, 16, 11, 10, 30))
    for m in r["aksiyon"]:
        _satir(ws, [m["standart"], m["madde"], m["bolum"], m["gereklilik"], m["kayit"], "Evet" if m["kritik"] else "", m["durum"], m["sorumlu"], m["termin"], m["kalan"], m["not"]],
               tarih_s=(9,), sar=(4, 5, 11), dolgu={6: KIRMIZI if m["kritik"] else None, 10: KIRMIZI if m["kalan"] is not None and m["kalan"] < 0 else None})
    ws = wb.create_sheet("Tüm Maddeler")
    _baslik(ws, ["Standart", "Bölüm", "Madde", "Gereklilik", "Gerekli Kayıt / Doküman", "Kritik", "Girilen Durum", "Değerlendirilen Durum", "Puan", "Son Güncelleme", "Sorumlu", "Termin"],
            (14, 20, 9, 36, 32, 7, 13, 13, 6, 12, 16, 11))
    renk = {"Hazır": YESIL, "Kısmen": SARI, "Güncellenmeli": SARI, "Eksik": KIRMIZI}
    for m in sorted(r["maddeler"], key=lambda m: (katla(m["standart"]), katla(m["bolum"]), katla(m["madde"]))):
        _satir(ws, [m["standart"], m["bolum"], m["madde"], m["gereklilik"], m["kayit"], "Evet" if m["kritik"] else "", m["girilen"], m["durum"], m["puan"], m["guncelleme"], m["sorumlu"],
                    m["termin"]], para_s=(9,), tarih_s=(10, 12), sar=(4, 5), dolgu={8: renk.get(m["durum"])})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, denetim: date, bugun: date, kismen: Decimal = Decimal(50), gecerlilik: int = 365) -> dict:
    maddeler, uy = oku(yol)
    r = analiz(maddeler, denetim, bugun, kismen, gecerlilik)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Denetim maddelerine göre doküman ve kayıt hazırlık durumunu takip eder.")
    p.add_argument("--liste", type=Path, help="Standart, Bölüm, Madde No, Gereklilik, Gerekli Kayıt / Doküman, Kritik, Sorumlu, Durum, Son Güncelleme, Termin, Not")
    p.add_argument("--denetim", help="Denetim tarihi GG.AA.YYYY (örnek veride 20.11.2026)")
    p.add_argument("--bugun", help="Rapor günü GG.AA.YYYY (varsayılan bugün; örnek veride 14.10.2026)")
    p.add_argument("--kismen", default="50", help="Kısmen / Güncellenmeli maddenin puan yüzdesi (varsayılan 50)")
    p.add_argument("--gecerlilik", type=int, default=365, help="Doküman / kayıt en fazla kaç günlük olabilir (varsayılan 365)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "denetim_hazirlik.xlsx")
    a = p.parse_args(argv)
    ornek = a.liste is None
    yol = a.liste or ORNEK / "denetim_listesi.csv"
    denetim = tarih(a.denetim) if a.denetim else (date(2026, 11, 20) if ornek else None)
    bugun = tarih(a.bugun) if a.bugun else (date(2026, 10, 14) if ornek else date.today())
    ks = para(a.kismen)
    if denetim is None or bugun is None or ks is None or not 0 <= ks <= 100 or a.gecerlilik < 1:
        print("[X] --denetim ve --bugun GG.AA.YYYY (kendi listenizde --denetim zorunlu), --kismen 0-100, --gecerlilik en az 1 olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, denetim, bugun, ks, a.gecerlilik)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    g = r["genel"]
    print(f"[OK] Denetime {r['kalan']} gün · {g['uyg']} uygulanabilir madde · hazırlık %{g['oran']} · {len(r['aksiyon'])} aksiyon · {g['kritik_eksik']} kritik madde hazır değil")
    for s in r["standart"]:
        print(f"     {s['ad']}: %{s['oran']} ({s['Hazır']} hazır / {s['uyg']})")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
