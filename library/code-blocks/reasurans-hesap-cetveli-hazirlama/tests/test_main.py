import json
import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "r.xlsx"
        cls.s = main.calistir(ORNEK / "trete.json", ORNEK / "prim_bordrosu.csv", ORNEK / "hasar_bordrosu.csv", cls.cikti)
        cls.h = {h.no: h for h in cls.s["hasarlar"]}
        cls.k = " ".join(a for _, a in cls.s["kontroller"])

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_cetvel(self):
        self.assertEqual(dict(self.s["alacak"]), {"Devredilen prim": D("530900.00"), "Prim depo iadesi": D(1200000), "Depo faizi (365 gün)": D("240000.00")})
        self.assertEqual(dict(self.s["borc"]), {"Reasürans komisyonu": D("159270.00"), "Devredilen ödenen hasar": D("3375000.00"),
                                                "Prim depo tutulan": D("185815.00")})
        self.assertEqual(self.s["bakiye"], D("-1749185.00"))
        self.assertEqual(sum(v for _, _, v in self.s["paylar"]), self.s["bakiye"])

    def test_devir_oranlari(self):
        self.assertEqual(self.h["HSR-101"].devir_odenen, D("750000.00"))      # 30M: 25/30
        self.assertEqual(self.h["HSR-102"].devir_odenen, D("2625000.00"))     # 40M: 35/40
        self.assertEqual(self.h["HSR-103"].oran, 0)                            # 4M saklama içinde
        self.assertEqual(self.h["HSR-104"].devir_muallak, D("875000.00"))     # 12M: 7/12
        p = {x.no + x.hareket: x for x in self.s["primler"]}
        self.assertEqual(p["YNG-0025Yeni"].oran, D(45) / D(60))               # kapasiteyle sınırlı
        self.assertEqual(p["YNG-0003İptal"].oran, D(10) / D(15))              # ilk poliçenin oranı
        self.assertEqual(p["YNG-0003İptal"].devir_prim, D("-4000.00"))

    def test_kontroller(self):
        for parca in ("YNG-0025: sigorta bedeli 60.000.000,00 TL", "10.000.000,00 TL için fakültatif", "HSR-102: devredilen hasar 3.062.500,00 TL nakit hasar",
                      "HSR-105: YNG-0099 numaralı poliçe prim bordrosunda yok", "HSR-106: hasar tarihi 15.08.2026 poliçe süresi dışında",
                      "HSR-107: ödeme tarihi 05.10.2026 dönem dışında", "YNG-0026: başlangıç/hareket tarihi 20.06.2026 dönem dışında"):
            self.assertIn(parca, self.k)
        self.assertEqual(self.h["HSR-106"].devir_odenen, 0)

    def test_kar_komisyonu(self):
        kk = self.s["kk"]
        self.assertEqual((kk["yonetim"], kk["kk"]), (D("26545.00"), 0))

    def test_excel(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Hesap Cetveli", "Prim Bordrosu", "Hasar Bordrosu", "Kontroller", "Parametreler"])
        bakiye = next(r for r in wb["Hesap Cetveli"].iter_rows(values_only=True) if r[0] == "BAKİYE")
        self.assertEqual(bakiye[1:5], ("reasürör öder", -1749185.0, -1049511.0, -699674.0))


class KotparTesti(unittest.TestCase):
    def test_elle_hesap(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t / "t.json").write_text(json.dumps({"ad": "QS", "tur": "kotpar", "devir_orani": 0.5, "komisyon_orani": 0.25, "prim_depo_orani": 0.4,
                                                  "donem_baslangic": "01.01.2026", "donem_bitis": "31.03.2026",
                                                  "reasurorler": [{"ad": "A", "pay": 0.7}, {"ad": "B", "pay": 0.2}]}), encoding="utf-8")
            (t / "p.csv").write_text("Poliçe No;Başlangıç;Bitiş;Sigorta Bedeli;Brüt Prim\nP1;10.01.2026;10.01.2027;1.000.000;10.000\n"
                                     "P2;01.02.2026;01.02.2027;2.000.000;0\n", encoding="utf-8")
            (t / "h.csv").write_text("Hasar No;Poliçe No;Hasar Tarihi;Ödeme Tarihi;Ödenen;Muallak\nH1;P1;01.03.2026;20.03.2026;3.000;1.000\n", encoding="utf-8")
            s = main.calistir(t / "t.json", t / "p.csv", t / "h.csv", t / "o.xlsx")
            # prim 5.000, komisyon 1.250, hasar 1.500, depo 2.000 → bakiye 250
            self.assertEqual(s["bakiye"], D("250.00"))
            self.assertEqual(s["muallak"], D("500.00"))
            k = " ".join(a for _, a in s["kontroller"])
            self.assertIn("P2: yeni poliçede prim sıfır", k)
            self.assertIn("Reasürör payları toplamı %90", k)

    def test_hatali_trete(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "t.json"
            y.write_text(json.dumps({"tur": "kotpar", "donem_baslangic": "01.01.2026", "donem_bitis": "31.03.2026"}), encoding="utf-8")
            with self.assertRaises(ValueError):
                main.trete_oku(y)
            self.assertEqual(main.main(["--trete", str(y), "--cikti", str(Path(tmp) / "x.xlsx")]), 1)


if __name__ == "__main__":
    unittest.main()
