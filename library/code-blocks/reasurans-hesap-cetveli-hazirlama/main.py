"""
Reasürans Hesap Cetveli Hazırlama — Workers / Workless kod bloğu
Sigorta › Reasürans › Reasürans Uzmanı

Orantılı (kotpar / eksedan) bir trete için dönem hesap cetvelini prim ve hasar bordrolarından hazırlar:
  - Poliçe bazında devir oranı: kotpar → sabit oran; eksedan → (sigorta bedeli − saklama payı) / sigorta bedeli,
    trete kapasitesiyle (satır sayısı × saklama payı) sınırlı. İptal / zeyil hareketleri ilk poliçenin oranını alır.
  - Devredilen prim, reasürans komisyonu, devredilen ödenen hasar, devredilen muallak (bilgi).
  - Prim depo: dönemde tutulan depo, önceki dönemden iade edilen depo ve depo faizi.
  - Bakiye ve reasürör paylarına dağılım; nakit hasar (cash call) limitini aşan hasarlar.
  - Kâr komisyonu ön hesabı (bilgi amaçlı; yıllık hesap ve zarar devri trete şartına göre yapılır).
  - Kontroller: kapasite aşımı (fakültatif gerekir), bordroda olmayan poliçeye hasar, poliçe süresi dışında hasar,
    eşi olmayan iptal, dönem dışı kayıtlar.
İnternete bağlanmaz.

Kullanım:
    python main.py                                                     # örnek eksedan trete, 2026 3. çeyrek
    python main.py --trete trete.json --primler prim_bordrosu.xlsx --hasarlar hasar_bordrosu.xlsx
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
SIFIR = Decimal(0)
KURUS = Decimal("0.01")
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")

PRIM_SUTUNLARI = {"no": ("police no", "police", "police numarasi"), "baslangic": ("baslangic", "baslangic tarihi", "tanzim tarihi"),
                  "bitis": ("bitis", "bitis tarihi"), "bedel": ("sigorta bedeli", "bedel"), "prim": ("brut prim", "prim"),
                  "hareket": ("hareket", "tip", "islem", "zeyil turu")}
HASAR_SUTUNLARI = {"no": ("hasar no", "hasar dosya no", "dosya no"), "police": ("police no", "police"), "hasar_tarihi": ("hasar tarihi",),
                   "odeme_tarihi": ("odeme tarihi",), "odenen": ("odenen", "odenen hasar", "odeme"), "muallak": ("muallak", "muallak hasar")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(str(x or "").strip(), f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif s.count(".") > 1 or (s.count(".") == 1 and len(s.split(".")[1]) == 3):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal) -> Decimal:
    return x.quantize(KURUS, ROUND_HALF_UP)


def tl(x) -> str:
    return "—" if x is None else f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


# ----------------------------------------------------------------------------
# Trete ve bordrolar
# ----------------------------------------------------------------------------

@dataclass
class Trete:
    ad: str
    tur: str
    komisyon_orani: Decimal
    donem_bas: date
    donem_bit: date
    reasurorler: list[tuple[str, Decimal]]
    devir_orani: Decimal | None = None
    saklama: Decimal | None = None
    satir: Decimal | None = None
    depo_orani: Decimal = SIFIR
    faiz_orani: Decimal = SIFIR
    nakit_hasar: Decimal | None = None
    kk_orani: Decimal | None = None
    yonetim_orani: Decimal = SIFIR
    onceki_depo: Decimal = SIFIR
    onceki_depo_tarihi: date | None = None

    @property
    def kapasite(self) -> Decimal | None:
        return self.saklama * self.satir if self.tur == "eksedan" else None


def trete_oku(yol: Path) -> Trete:
    j = json.loads(yol.read_text(encoding="utf-8-sig"))
    d = lambda k, v=None: Decimal(str(j[k])) if j.get(k) is not None else v  # noqa: E731
    tur = katla(j.get("tur"))
    tur = "kotpar" if tur in ("kotpar", "quota share", "qs") else "eksedan" if tur in ("eksedan", "surplus") else tur
    if tur not in ("kotpar", "eksedan"):
        raise ValueError("trete 'tur' kotpar veya eksedan olmalı")
    t = Trete(j.get("ad", "Trete"), tur, d("komisyon_orani", SIFIR), tarih(j["donem_baslangic"]), tarih(j["donem_bitis"]),
              [(r["ad"], Decimal(str(r["pay"]))) for r in j.get("reasurorler", [])], d("devir_orani"), d("saklama_payi"), d("satir_sayisi"),
              d("prim_depo_orani", SIFIR), d("depo_faiz_orani", SIFIR), d("nakit_hasar_limiti"), d("kar_komisyonu_orani"), d("yonetim_gideri_orani", SIFIR))
    od = j.get("onceki_depo") or {}
    t.onceki_depo, t.onceki_depo_tarihi = Decimal(str(od.get("tutar", 0))), tarih(od.get("tutulma_tarihi"))
    if tur == "kotpar" and t.devir_orani is None:
        raise ValueError("kotpar trete için 'devir_orani' gerekli")
    if tur == "eksedan" and (t.saklama is None or t.satir is None):
        raise ValueError("eksedan trete için 'saklama_payi' ve 'satir_sayisi' gerekli")
    if t.donem_bas is None or t.donem_bit is None:
        raise ValueError("donem_baslangic ve donem_bitis GG.AA.YYYY olmalı")
    return t


@dataclass
class Prim:
    satir: int
    no: str
    baslangic: date | None
    bitis: date | None
    bedel: Decimal
    prim: Decimal
    hareket: str
    oran: Decimal = SIFIR
    devir_bedeli: Decimal = SIFIR
    devir_prim: Decimal = SIFIR
    komisyon: Decimal = SIFIR
    notlar: list = field(default_factory=list)


@dataclass
class Hasar:
    satir: int
    no: str
    police: str
    hasar_tarihi: date | None
    odeme_tarihi: date | None
    odenen: Decimal
    muallak: Decimal
    oran: Decimal = SIFIR
    devir_odenen: Decimal = SIFIR
    devir_muallak: Decimal = SIFIR
    notlar: list = field(default_factory=list)


def devir_orani(t: Trete, bedel: Decimal) -> tuple[Decimal, Decimal]:
    """(oran, devredilen bedel)."""
    if t.tur == "kotpar":
        return t.devir_orani, bedel * t.devir_orani
    if bedel <= t.saklama:
        return SIFIR, SIFIR
    devir = min(bedel - t.saklama, t.kapasite)
    return devir / bedel, devir


# ----------------------------------------------------------------------------
# Hesap
# ----------------------------------------------------------------------------

def hesapla(t: Trete, primler_ham: list[dict], hasarlar_ham: list[dict]) -> dict:
    kontroller = []

    def k(onem, aciklama):
        kontroller.append((onem, aciklama))

    def donemde(d):
        return d is not None and t.donem_bas <= d <= t.donem_bit

    primler, ilk = [], {}
    for r in primler_ham:
        if not r.get("no"):
            continue
        p = Prim(r["_satir"], str(r["no"]).strip(), tarih(r.get("baslangic")), tarih(r.get("bitis")), para(r.get("bedel")) or SIFIR,
                 para(r.get("prim")) or SIFIR, str(r.get("hareket") or "Yeni").strip())
        primler.append(p)
    for p in primler:
        hk = katla(p.hareket)
        if not any(x in hk for x in ("iptal", "zeyil", "iade", "ek prim")) and p.no not in ilk:
            ilk[p.no] = p
    dahil = []
    for p in primler:
        if not donemde(p.baslangic):
            p.notlar.append("Dönem dışı — hesaba alınmadı")
            k("Bilgi", f"{p.no}: başlangıç/hareket tarihi {p.baslangic:%d.%m.%Y} dönem dışında, hesaba alınmadı" if p.baslangic else f"{p.no}: tarih okunamadı")
            continue
        asil = ilk.get(p.no)
        if asil is not p:
            if asil is None:
                k("Orta", f"{p.no}: {p.hareket} hareketinin ilk poliçe kaydı bordroda yok; devir oranı bilinmiyor, hesaba alınmadı")
                p.notlar.append("İlk poliçe kaydı yok")
                continue
            p.oran = asil.oran if asil.oran or asil in dahil else devir_orani(t, asil.bedel)[0]
            p.notlar.append(f"Oran ilk poliçeden ({asil.no})")
        else:
            p.oran, p.devir_bedeli = devir_orani(t, p.bedel)
            if t.tur == "eksedan" and p.bedel > t.saklama + t.kapasite:
                fazla = p.bedel - t.saklama - t.kapasite
                p.notlar.append(f"Kapasite aşımı {tl(fazla)} TL")
                k("Yüksek", f"{p.no}: sigorta bedeli {tl(p.bedel)} TL, saklama + trete kapasitesi {tl(t.saklama + t.kapasite)} TL; "
                            f"{tl(fazla)} TL için fakültatif plasman gerekir")
            if p.prim <= 0:
                k("Orta", f"{p.no}: yeni poliçede prim sıfır veya negatif")
        p.devir_prim = yuvarla(p.prim * p.oran)
        p.komisyon = yuvarla(p.devir_prim * t.komisyon_orani)
        dahil.append(p)

    police = {p.no: p for p in primler if ilk.get(p.no) is p}
    hasarlar = []
    for r in hasarlar_ham:
        if not r.get("no"):
            continue
        h = Hasar(r["_satir"], str(r["no"]).strip(), str(r.get("police") or "").strip(), tarih(r.get("hasar_tarihi")), tarih(r.get("odeme_tarihi")),
                  para(r.get("odenen")) or SIFIR, para(r.get("muallak")) or SIFIR)
        hasarlar.append(h)
        p = police.get(h.police)
        if p is None:
            k("Yüksek", f"{h.no}: {h.police} numaralı poliçe prim bordrosunda yok; devir oranı bilinmiyor, hesaba alınmadı")
            h.notlar.append("Poliçe bordroda yok")
            continue
        h.oran = p.oran if p.oran or p in dahil else devir_orani(t, p.bedel)[0]
        if h.hasar_tarihi and p.baslangic and p.bitis and not (p.baslangic <= h.hasar_tarihi < p.bitis):
            k("Yüksek", f"{h.no}: hasar tarihi {h.hasar_tarihi:%d.%m.%Y} poliçe süresi dışında ({p.baslangic:%d.%m.%Y} – {p.bitis:%d.%m.%Y}); "
                        "hesaba alınmadı")
            h.notlar.append("Hasar tarihi poliçe süresi dışında — hesaba alınmadı")
            continue
        if h.odenen and not donemde(h.odeme_tarihi):
            h.notlar.append("Ödeme dönem dışında — sonraki cetvele")
            k("Bilgi", f"{h.no}: ödeme tarihi {h.odeme_tarihi:%d.%m.%Y} dönem dışında; bu cetvele alınmadı" if h.odeme_tarihi else f"{h.no}: ödeme tarihi yok")
        else:
            h.devir_odenen = yuvarla(h.odenen * h.oran)
        h.devir_muallak = yuvarla(h.muallak * h.oran)
        if not h.oran and (h.odenen or h.muallak):
            h.notlar.append("Saklama payı içinde — devir yok")
        if t.nakit_hasar is not None and h.devir_odenen + h.devir_muallak >= t.nakit_hasar:
            k("Orta", f"{h.no}: devredilen hasar {tl(h.devir_odenen + h.devir_muallak)} TL nakit hasar limitini ({tl(t.nakit_hasar)} TL) aşıyor — "
                      "reasürörlerden nakit hasar (cash call) talep edilebilir")
            h.notlar.append("Nakit hasar limiti aşıldı")

    prim = sum((p.devir_prim for p in dahil), SIFIR)
    komisyon = sum((p.komisyon for p in dahil), SIFIR)
    odenen = sum((h.devir_odenen for h in hasarlar), SIFIR)
    muallak = sum((h.devir_muallak for h in hasarlar), SIFIR)
    depo_tutulan = yuvarla(prim * t.depo_orani)
    depo_iade, faiz, faiz_gun = SIFIR, SIFIR, 0
    if t.onceki_depo and t.onceki_depo_tarihi:
        depo_iade = t.onceki_depo
        faiz_gun = (t.donem_bit - t.onceki_depo_tarihi).days
        faiz = yuvarla(t.onceki_depo * t.faiz_orani * faiz_gun / 365)
    alacak = [("Devredilen prim", prim), ("Prim depo iadesi", depo_iade), (f"Depo faizi ({faiz_gun} gün)", faiz)]
    borc = [("Reasürans komisyonu", komisyon), ("Devredilen ödenen hasar", odenen), ("Prim depo tutulan", depo_tutulan)]
    bakiye = sum(v for _, v in alacak) - sum(v for _, v in borc)
    toplam_pay = sum(p for _, p in t.reasurorler)
    if t.reasurorler and toplam_pay != 1:
        k("Yüksek", f"Reasürör payları toplamı %{toplam_pay * 100:g}; %100 olmalı")
    kk = None
    if t.kk_orani is not None:
        yonetim = yuvarla(prim * t.yonetim_orani)
        kar = prim - komisyon - odenen - muallak - yonetim
        kk = {"prim": prim, "komisyon": komisyon, "odenen": odenen, "muallak": muallak, "yonetim": yonetim, "kar": kar,
              "kk": yuvarla(max(kar, SIFIR) * t.kk_orani)}
    return {"primler": primler, "dahil": dahil, "hasarlar": hasarlar, "kontroller": kontroller, "alacak": alacak, "borc": borc, "bakiye": bakiye,
            "muallak": muallak, "kk": kk, "paylar": paylara_bol(bakiye, t.reasurorler)}


def paylara_bol(tutar: Decimal, reasurorler: list[tuple[str, Decimal]]) -> list[tuple[str, Decimal, Decimal]]:
    """Kuruş farkı son reasüröre yazılır; paylar toplamı %100 ise dağılım toplamı bakiyeye eşittir."""
    sonuc = [(ad, pay, yuvarla(tutar * pay)) for ad, pay in reasurorler]
    if sonuc and sum(p for _, p in reasurorler) == 1:
        fark = tutar - sum(v for _, _, v in sonuc)
        ad, pay, v = sonuc[-1]
        sonuc[-1] = (ad, pay, v + fark)
    return sonuc


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KONTROL = PatternFill("solid", fgColor="FFF4CE")
UST = Alignment(vertical="top", wrap_text=True)
SAYI = "#,##0.00"


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def rapor_yaz(cikti: Path, t: Trete, s: dict) -> None:
    wb = Workbook()
    hc = wb.active
    hc.title = "Hesap Cetveli"
    adlar = [ad for ad, _ in t.reasurorler]
    hc.column_dimensions["A"].width = 36
    for j in range(2, 4 + len(adlar)):
        hc.column_dimensions[get_column_letter(j)].width = 18
    hc.append([t.ad])
    hc.cell(1, 1).font = Font(bold=True, size=13)
    hc.append([f"Dönem: {t.donem_bas:%d.%m.%Y} – {t.donem_bit:%d.%m.%Y} · Tür: {t.tur} · Komisyon %{t.komisyon_orani * 100:g} · "
               f"Prim depo %{t.depo_orani * 100:g} · Depo faizi %{t.faiz_orani * 100:g}"])
    hc.append([])
    hc.append(["Kalem", "Taraf", "%100", *[f"{ad} (%{p * 100:g})" for ad, p in t.reasurorler]])
    for h in hc[hc.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for taraf, kalemler in (("Alacak (reasürör lehine)", s["alacak"]), ("Borç (reasürör aleyhine)", s["borc"])):
        for a, v in kalemler:
            hc.append([a, taraf, float(v), *[float(yuvarla(v * p)) for _, p in t.reasurorler]])
            for j in range(3, 4 + len(adlar)):
                hc.cell(hc.max_row, j).number_format = SAYI
    hc.append(["BAKİYE", "şirket öder" if s["bakiye"] >= 0 else "reasürör öder", float(s["bakiye"]), *[float(v) for _, _, v in s["paylar"]]])
    for j in range(1, 4 + len(adlar)):
        hc.cell(hc.max_row, j).font = Font(bold=True)
        if j >= 3:
            hc.cell(hc.max_row, j).number_format = SAYI
    hc.append([])
    hc.append(["Devredilen muallak hasar (bilgi)", "", float(s["muallak"]), *[float(yuvarla(s["muallak"] * p)) for _, p in t.reasurorler]])
    for j in range(3, 4 + len(adlar)):
        hc.cell(hc.max_row, j).number_format = SAYI
    if s["kk"]:
        hc.append([])
        hc.append(["Kâr komisyonu ön hesabı (bilgi)"])
        hc.cell(hc.max_row, 1).font = Font(bold=True)
        for a, key in (("Devredilen prim", "prim"), ("− Komisyon", "komisyon"), ("− Ödenen hasar", "odenen"), ("− Muallak hasar", "muallak"),
                       (f"− Yönetim gideri (%{t.yonetim_orani * 100:g})", "yonetim"), ("= Kâr", "kar"), (f"Kâr komisyonu (%{t.kk_orani * 100:g})", "kk")):
            hc.append([a, "", float(s["kk"][key])])
            hc.cell(hc.max_row, 3).number_format = SAYI
        hc.append(["Not: Kâr komisyonu yıllık ve trete şartına (zarar devri, muallak tanımı) göre hesaplanır; bu satırlar yalnız fikir verir."])

    pb = wb.create_sheet("Prim Bordrosu")
    _baslik(pb, ["Poliçe No", "Başlangıç", "Bitiş", "Sigorta Bedeli", "Brüt Prim", "Hareket", "Devir Oranı", "Devredilen Bedel", "Devredilen Prim",
                 "Komisyon", "Not"], (11, 11, 11, 15, 12, 9, 10, 15, 14, 12, 40))
    for p in s["primler"]:
        pb.append([p.no, p.baslangic, p.bitis, float(p.bedel), float(p.prim), p.hareket, float(p.oran), float(p.devir_bedeli), float(p.devir_prim),
                   float(p.komisyon), "; ".join(p.notlar)])
        for c in (2, 3):
            pb.cell(pb.max_row, c).number_format = "DD.MM.YYYY"
        for c in (4, 5, 8, 9, 10):
            pb.cell(pb.max_row, c).number_format = SAYI
        pb.cell(pb.max_row, 7).number_format = "0.00%"
        if p.notlar:
            pb.cell(pb.max_row, 11).fill = KONTROL
    pb.append(["TOPLAM", None, None, None, float(sum((p.prim for p in s["dahil"]), SIFIR)), None, None, None,
               float(sum((p.devir_prim for p in s["dahil"]), SIFIR)), float(sum((p.komisyon for p in s["dahil"]), SIFIR))])
    for c in (5, 9, 10):
        pb.cell(pb.max_row, c).number_format = SAYI
    for h in pb[pb.max_row]:
        h.font = Font(bold=True)

    hb = wb.create_sheet("Hasar Bordrosu")
    _baslik(hb, ["Hasar No", "Poliçe No", "Hasar Tarihi", "Ödeme Tarihi", "Ödenen", "Muallak", "Devir Oranı", "Devredilen Ödenen", "Devredilen Muallak", "Not"],
            (10, 11, 11, 11, 13, 13, 10, 15, 15, 44))
    for h in s["hasarlar"]:
        hb.append([h.no, h.police, h.hasar_tarihi, h.odeme_tarihi, float(h.odenen), float(h.muallak), float(h.oran), float(h.devir_odenen),
                   float(h.devir_muallak), "; ".join(h.notlar)])
        for c in (3, 4):
            hb.cell(hb.max_row, c).number_format = "DD.MM.YYYY"
        for c in (5, 6, 8, 9):
            hb.cell(hb.max_row, c).number_format = SAYI
        hb.cell(hb.max_row, 7).number_format = "0.00%"
        if h.notlar:
            hb.cell(hb.max_row, 10).fill = KONTROL

    kt = wb.create_sheet("Kontroller")
    _baslik(kt, ["Önem", "Kontrol", "İnceleme"], (9, 110, 26))
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    for o, a in sorted(s["kontroller"], key=lambda x: sira[x[0]]):
        kt.append([o, a, ""])
        kt.cell(kt.max_row, 1).fill = PatternFill("solid", fgColor=RENK[o])
        kt.cell(kt.max_row, 3).fill = KONTROL
        kt.cell(kt.max_row, 2).alignment = UST

    pa = wb.create_sheet("Parametreler")
    _baslik(pa, ["Parametre", "Değer"], (34, 50))
    for a, d in [("Trete", t.ad), ("Tür", t.tur), ("Devir oranı (kotpar)", t.devir_orani), ("Saklama payı (eksedan)", t.saklama), ("Satır sayısı", t.satir),
                 ("Trete kapasitesi", t.kapasite), ("Komisyon oranı", t.komisyon_orani), ("Prim depo oranı", t.depo_orani), ("Depo faiz oranı (yıllık)", t.faiz_orani),
                 ("Nakit hasar limiti", t.nakit_hasar), ("Önceki dönem depo", t.onceki_depo),
                 ("Önceki depo tutulma tarihi", f"{t.onceki_depo_tarihi:%d.%m.%Y}" if t.onceki_depo_tarihi else None)]:
        pa.append([a, float(d) if isinstance(d, Decimal) else d])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(trete_yolu: Path, prim_yolu: Path, hasar_yolu: Path | None, cikti: Path) -> dict:
    t = trete_oku(trete_yolu)
    primler = kayitlar(prim_yolu, PRIM_SUTUNLARI, ("no", "prim"))
    hasarlar = kayitlar(hasar_yolu, HASAR_SUTUNLARI, ("no", "police")) if hasar_yolu else []
    s = hesapla(t, primler, hasarlar)
    rapor_yaz(cikti, t, s)
    return {**s, "trete": t}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Orantılı trete için prim ve hasar bordrolarından reasürans hesap cetveli hazırlar.")
    p.add_argument("--trete", type=Path, default=ORNEK / "trete.json", help="Trete şartları (.json)")
    p.add_argument("--primler", type=Path, default=ORNEK / "prim_bordrosu.csv", help="Prim bordrosu (.xlsx/.csv)")
    p.add_argument("--hasarlar", type=Path, default=None, help="Hasar bordrosu (.xlsx/.csv)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "reasurans_hesap_cetveli.xlsx")
    a = p.parse_args(argv)
    hasarlar = a.hasarlar or (ORNEK / "hasar_bordrosu.csv" if a.primler == ORNEK / "prim_bordrosu.csv" else None)
    for y in (a.trete, a.primler, hasarlar):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        s = calistir(a.trete, a.primler, hasarlar, a.cikti)
    except (ValueError, KeyError, json.JSONDecodeError) as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {s['trete'].ad}: {len(s['dahil'])} prim hareketi, {len(s['hasarlar'])} hasar")
    print(f"[OK] Bakiye {tl(s['bakiye'])} TL ({'şirket öder' if s['bakiye'] >= 0 else 'reasürör öder'}) · " +
          " · ".join(f"{ad} {tl(v)}" for ad, _, v in s["paylar"]))
    for o, k in s["kontroller"]:
        if o != "Bilgi":
            print(f"[{'X' if o == 'Yüksek' else '!'}] {k}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
