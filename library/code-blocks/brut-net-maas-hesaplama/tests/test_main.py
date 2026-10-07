"""Beklenen değerler resmî/bağımsız kaynaklardan alınmıştır (README'de listelenmiştir)."""
import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))

import bordro  # noqa: E402
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

P = bordro.Parametreler(2026)


class AsgariUcretTesti(unittest.TestCase):
    def test_net_ve_isveren_maliyeti(self):
        b = bordro.ay_hesapla(P, 1, D("33030"))
        self.assertEqual(b.net, D("28075.50"))                        # Resmî Gazete 26.12.2025
        self.assertEqual((b.sgk_isci, b.issizlik_isci), (D("4624.20"), D("330.30")))
        self.assertEqual((b.odenecek_gv, b.odenecek_dv), (D("0.00"), D("0.00")))
        self.assertEqual(b.isveren_maliyeti, D("40874.63"))             # teşviksiz
        self.assertEqual(bordro.ay_hesapla(P, 1, D("33030"), tesvik="imalat").isveren_maliyeti, D("39223.13"))
        self.assertEqual(bordro.ay_hesapla(P, 1, D("33030"), tesvik="diger").isveren_maliyeti, D("40214.03"))

    def test_yillik_gv_istisnasi(self):
        yil = bordro.yil_hesapla(P, [(m, D("33030"), 30) for m in range(1, 13)])
        aylik = [b.gv_istisnasi for b in yil]
        self.assertEqual(aylik[0], D("4211.33"))
        self.assertEqual(aylik[6], D("4537.75"))                       # Temmuz: dilim geçişi
        self.assertEqual(aylik[7], D("5615.10"))
        self.assertEqual(sum(aylik), D("57881.23"))                    # yıllık toplam istisna

    def test_eksik_gun(self):
        b = bordro.ay_hesapla(P, 1, D("16515"), gun=15)                # 15 günlük asgari ücret
        self.assertEqual((b.odenecek_gv, b.odenecek_dv), (D("0.00"), D("0.00")))


class UcretTesti(unittest.TestCase):
    def test_ocak_50000(self):
        b = bordro.ay_hesapla(P, 1, D("50000"))
        self.assertEqual(b.odenecek_gv, D("2163.67"))
        self.assertEqual(b.odenecek_dv, D("128.80"))
        self.assertEqual(b.net, D("40207.53"))

    def test_dilim_gecisi_mart_100000(self):
        yil = bordro.yil_hesapla(P, [(m, D("100000"), 30) for m in range(1, 13)])
        # Mart: kümülatif 170.000 → 255.000; vergi 28.500 + 65.000×%20 − 25.500 = 16.000; − 4.211,33 istisna
        self.assertEqual(yil[2].hesaplanan_gv, D("16000.00"))
        self.assertEqual(yil[2].net, D("72703.03"))
        self.assertEqual(yil[2].vergi_dilimi, D("20.00"))

    def test_sgk_tavani(self):
        b = bordro.ay_hesapla(P, 1, D("400000"))
        self.assertEqual(b.sgk_matrahi, D("297270.00"))
        self.assertEqual(b.sgk_isci, D("41617.80"))
        self.assertIn("tavan", b.uyarilar)

    def test_netten_brute_tersi(self):
        for ay, net in ((1, D("40207.53")), (9, D("60000")), (12, D("150000"))):
            b = bordro.netten_brute(P, ay, net, D(0) if ay == 1 else D("500000"))
            self.assertEqual(b.net, net)
            # bir kuruş daha az brüt hedef nete ulaşamaz
            self.assertLess(bordro.ay_hesapla(P, ay, b.brut - D("0.01"), D(0) if ay == 1 else D("500000")).net, net)


class UctanUcaTest(unittest.TestCase):
    def test_ornek_personel(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "b.xlsx"
            sonuc = main.calistir(P, cikti, girdi=KLASOR / "ornek_veri" / "personel.csv")
            self.assertEqual(set(sonuc), {"Ayşe Yılmaz", "Mehmet Kaya", "Zeynep Demir", "Ali Veli"})
            self.assertTrue(all(b.net == D("60000.00") for b in sonuc["Mehmet Kaya"]))
            self.assertEqual(len(sonuc["Ali Veli"]), 9)                 # Nisan-Aralık
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames[0], "Özet")
            self.assertIn("Parametreler", wb.sheetnames)


if __name__ == "__main__":
    unittest.main()
