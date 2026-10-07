"""
Brüt-Net Maaş Hesaplama — Workers / Workless kod bloğu
İnsan Kaynakları › Bordro ve Özlük İşleri Uzmanı (ayrıca: Mali Müşavirlik › Bordro Elemanı)

Türkiye'nin güncel yıl parametreleriyle (SGK, işsizlik, gelir vergisi tarifesi, asgari ücret
gelir ve damga vergisi istisnası) brütten nete ve netten brüte aylık bordro hesaplar; 12 aylık
tabloyu, işveren maliyetini ve vergi dilimi geçişlerini Excel'e yazar. İnternete bağlanmaz.

Kullanım:
    python main.py --brut 85000                       # 12 ay, her ay 85.000 TL brüt
    python main.py --net 60000                        # her ay 60.000 TL net için gereken brütler
    python main.py --girdi personel.xlsx              # kişi bazında toplu hesap
    python main.py --brut 85000 --tesvik imalat --yil 2026
"""
from __future__ import annotations

import argparse
import csv
import sys
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import bordro

BURASI = Path(__file__).resolve().parent

SUTUNLAR = [
    ("Ay", "ay"), ("Gün", "gun"), ("Brüt", "brut"), ("SGK Matrahı", "sgk_matrahi"), ("SGK İşçi", "sgk_isci"),
    ("İşsizlik İşçi", "issizlik_isci"), ("GV Matrahı", "gv_matrahi"), ("Kümülatif GV Matrahı", "kumulatif_gv_matrahi"),
    ("Vergi Dilimi %", "vergi_dilimi"), ("Hesaplanan GV", "hesaplanan_gv"), ("AÜ GV İstisnası", "gv_istisnasi"),
    ("Ödenecek GV", "odenecek_gv"), ("Hesaplanan DV", "hesaplanan_dv"), ("AÜ DV İstisnası", "dv_istisnasi"),
    ("Ödenecek DV", "odenecek_dv"), ("Net Ücret", "net"), ("SGK İşveren", "sgk_isveren"),
    ("İşsizlik İşveren", "issizlik_isveren"), ("İşveren Maliyeti", "isveren_maliyeti"), ("Uyarılar", "uyarilar"),
]
AYLAR = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
TOPLAM_YAZI = Font(bold=True)
PARA = "#,##0.00"


def sayi(x) -> Decimal | None:
    if x is None or str(x).strip() == "":
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        raise ValueError(f"Sayı anlaşılamadı: {x!r}")


def tl(x) -> str:
    return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def tablo_yaz(ws, satirlar: list[bordro.AyBordrosu], baslangic_satir: int = 1) -> int:
    for j, (baslik, _) in enumerate(SUTUNLAR, 1):
        h = ws.cell(baslangic_satir, j, baslik)
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    r = baslangic_satir
    for b in satirlar:
        r += 1
        d = b.sozluk()
        for j, (_, alan) in enumerate(SUTUNLAR, 1):
            v = d[alan]
            if alan == "ay":
                v = AYLAR[v - 1]
            h = ws.cell(r, j, float(v) if isinstance(v, Decimal) else v)
            if isinstance(v, Decimal) and alan != "vergi_dilimi":
                h.number_format = PARA
    r += 1
    ws.cell(r, 1, "TOPLAM").font = TOPLAM_YAZI
    for j, (_, alan) in enumerate(SUTUNLAR, 1):
        if alan in {"ay", "gun", "vergi_dilimi", "kumulatif_gv_matrahi", "uyarilar"}:
            continue
        h = ws.cell(r, j, float(sum((getattr(b, alan) for b in satirlar), Decimal(0))))
        h.number_format, h.font = PARA, TOPLAM_YAZI
    for j, (baslik, _) in enumerate(SUTUNLAR, 1):
        ws.column_dimensions[get_column_letter(j)].width = 40 if baslik == "Uyarılar" else max(11, len(baslik) + 2)
    return r


def personel_oku(yol: Path) -> list[dict]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ayirici = csv.Sniffer().sniff(metin.splitlines()[0], delimiters=",;\t").delimiter
        satirlar = list(csv.reader(metin.splitlines(), delimiter=ayirici))
    satirlar = [r for r in satirlar if any(c not in (None, "") for c in r)]
    # Türkçe küçük harf: Python'un lower() 'İ'yi 'i̇' (birleşik nokta) yapar
    basliklar = [str(b or "").replace("İ", "i").replace("I", "ı").lower().strip() for b in satirlar[0]]

    def bul(*adlar):
        for i, b in enumerate(basliklar):
            if b in adlar:
                return i
        return None

    i_ad = bul("ad soyad", "ad", "personel", "çalışan", "calisan", "sicil")
    i_brut, i_net = bul("brüt", "brut", "brüt ücret", "brut ucret"), bul("net", "net ücret", "net ucret")
    i_gun, i_kum = bul("gün", "gun"), bul("önceki kümülatif matrah", "onceki kumulatif matrah", "kümülatif matrah")
    i_bas = bul("başlangıç ayı", "baslangic ayi", "ay")
    if i_ad is None or (i_brut is None and i_net is None):
        raise SystemExit(f"Başlıklar tanınmadı. Gerekli: 'Ad Soyad' ve 'Brüt' veya 'Net'. Bulunan: {satirlar[0]}")
    kisiler = []
    for r in satirlar[1:]:
        brut = sayi(r[i_brut]) if i_brut is not None else None
        net = sayi(r[i_net]) if i_net is not None else None
        kisiler.append({
            "ad": str(r[i_ad]).strip(),
            "tutar": brut if brut is not None else net,
            "net_hedefli": brut is None,
            "gun": int(sayi(r[i_gun]) or 30) if i_gun is not None else 30,
            "kumulatif": (sayi(r[i_kum]) or Decimal(0)) if i_kum is not None else Decimal(0),
            "baslangic": int(sayi(r[i_bas]) or 1) if i_bas is not None else 1,
        })
    return kisiler


def calistir(p: bordro.Parametreler, cikti: Path, *, brut=None, net=None, girdi: Path | None = None,
             tesvik="yok", baslangic_ayi=1) -> dict:
    wb = Workbook()
    wb.remove(wb.active)
    sonuc = {}
    if girdi:
        kisiler = personel_oku(girdi)
        ozet = wb.create_sheet("Özet")
        ozet.append(["Ad Soyad", "Hesap Türü", "Yıllık Brüt", "Yıllık Net", "Yıllık İşveren Maliyeti", "En Düşük Net", "En Yüksek Net"])
        for h in ozet[1]:
            h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        for k in kisiler:
            aylar = [(m, k["tutar"], k["gun"]) for m in range(k["baslangic"], 13)]
            tablo = bordro.yil_hesapla(p, aylar, tesvik, k["kumulatif"], k["net_hedefli"])
            ws = wb.create_sheet(k["ad"][:28] or "Personel")
            tablo_yaz(ws, tablo)
            ozet.append([k["ad"], "Netten brüte" if k["net_hedefli"] else "Brütten nete",
                         float(sum(b.brut for b in tablo)), float(sum(b.net for b in tablo)),
                         float(sum(b.isveren_maliyeti for b in tablo)),
                         float(min(b.net for b in tablo)), float(max(b.net for b in tablo))])
            sonuc[k["ad"]] = tablo
        for j in range(1, 8):
            ozet.column_dimensions[get_column_letter(j)].width = 20
            for h in ozet[get_column_letter(j)][1:]:
                if isinstance(h.value, float):
                    h.number_format = PARA
    else:
        tutar, net_hedefli = (Decimal(str(brut)), False) if brut is not None else (Decimal(str(net)), True)
        tablo = bordro.yil_hesapla(p, [(m, tutar, 30) for m in range(baslangic_ayi, 13)], tesvik, Decimal(0), net_hedefli)
        ws = wb.create_sheet("Bordro")
        tablo_yaz(ws, tablo)
        sonuc["Bordro"] = tablo

    bilgi = wb.create_sheet("Parametreler")
    bilgi.append(["Yıl", p.yil])
    bilgi.append(["İşveren teşviki", tesvik])
    for anahtar in ("asgari_ucret_brut", "sgk_isci_orani", "issizlik_isci_orani", "sgk_isveren_orani",
                    "issizlik_isveren_orani", "sgk_tavan", "damga_vergisi_orani"):
        bilgi.append([anahtar, p.ham[anahtar]])
    bilgi.append(["gelir_vergisi_ucret_tarifesi", "; ".join(f"{s or 'üstü'} → %{Decimal(o) * 100:g}" for s, o in p.ham["gelir_vergisi_ucret_tarifesi"])])
    for k in p.ham.get("kaynaklar", []):
        bilgi.append(["Kaynak", k])
    bilgi.append(["Not", "Hesaplama bilgilendirme amaçlıdır; engellilik indirimi, BES kesintisi, ayni yardımlar ve özel teşvikler dahil değildir."])
    bilgi.column_dimensions["A"].width = 30
    bilgi.column_dimensions["B"].width = 120

    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    return sonuc


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description="Brütten nete / netten brüte aylık bordro hesabı (Türkiye).")
    g = ap.add_mutually_exclusive_group()
    g.add_argument("--brut", type=sayi, help="Aylık brüt ücret")
    g.add_argument("--net", type=sayi, help="Aylık hedef net ücret")
    g.add_argument("--girdi", type=Path, help="Personel listesi (.xlsx/.csv): Ad Soyad + Brüt veya Net [+ Gün, Önceki Kümülatif Matrah, Başlangıç Ayı]")
    ap.add_argument("--yil", type=int, default=2026)
    ap.add_argument("--tesvik", choices=["yok", "imalat", "diger"], default="yok", help="İşveren SGK prim teşviki")
    ap.add_argument("--baslangic-ayi", type=int, default=1, help="Tek kişilik hesapta ilk ay (1-12)")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "bordro.xlsx")
    a = ap.parse_args(argv)
    if a.brut is None and a.net is None and a.girdi is None:
        a.girdi = BURASI / "ornek_veri" / "personel.csv"

    p = bordro.Parametreler(a.yil)
    sonuc = calistir(p, a.cikti, brut=a.brut, net=a.net, girdi=a.girdi, tesvik=a.tesvik, baslangic_ayi=a.baslangic_ayi)
    for ad, tablo in sonuc.items():
        ilk, son = tablo[0], tablo[-1]
        print(f"[OK] {ad}: {AYLAR[ilk.ay - 1]} brüt {tl(ilk.brut)} → net {tl(ilk.net)} · "
              f"{AYLAR[son.ay - 1]} brüt {tl(son.brut)} → net {tl(son.net)}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
