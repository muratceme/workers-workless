"""
Kritik Stok Seviyesi Raporu — Workers / Workless kod bloğu
Eczane ve Tıbbi Malzeme › Hastane Eczacısı
Günlük tüketim ve tedarik süresine göre kritik stok seviyesinin altına düşen ilaç ve malzemeleri listeler.
  Günlük tüketim     = Günlük Tüketim sütunu · ya da Aylık Tüketim ÷ 30 · ya da --tuketim dosyasındaki dönem tüketimi ÷ --gun
  Emniyet stoğu      = günlük tüketim × emniyet günü (satırda yoksa --emniyet, 7)
  Kritik seviye      = günlük tüketim × (tedarik süresi + emniyet günü)   — sipariş verilmesi gereken seviye
  Stok günü          = mevcut stok ÷ günlük tüketim; tahmini tükenme = rapor tarihi + stok günü
  Önerilen sipariş   = günlük tüketim × (tedarik süresi + emniyet günü + --kapsam (30)) − mevcut stok − bekleyen sipariş (yukarı yuvarlanır);
                       yalnız Stoksuz ve Kritik altında malzemeler için
  Durum              = Stoksuz · Kritik altında (mevcut + bekleyen < kritik seviye) · Sipariş yolda (mevcut < kritik seviye ≤ mevcut + bekleyen) · Yeterli
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

STOK = {"kod": ("malzeme kodu", "ilac kodu", "stok kodu", "barkod", "kod"), "ad": ("malzeme adi", "ilac adi", "urun adi", "malzeme", "ad"), "birim": ("birim", "olcu birimi"),
        "mevcut": ("mevcut stok", "stok", "eldeki stok", "stok miktari", "miktar"), "gunluk": ("gunluk tuketim", "gunluk ortalama tuketim"), "aylik": ("aylik tuketim", "aylik ortalama tuketim"),
        "tedarik": ("tedarik suresi", "temin suresi", "teslim suresi", "tedarik suresi gun"), "emniyet": ("emniyet gunu", "emniyet suresi", "guvenlik gunu"),
        "bekleyen": ("bekleyen siparis", "yoldaki siparis", "acik siparis"), "fiyat": ("birim fiyat", "birim maliyet", "fiyat"), "hayati": ("hayati", "kritik ilac", "hayati onemde", "kritik")}
TUKETIM = {"kod": ("malzeme kodu", "ilac kodu", "stok kodu", "barkod", "kod"), "miktar": ("tuketim", "tuketim miktari", "miktar", "cikis miktari")}
SIFIR = Decimal(0)
EVET = {"evet", "e", "x", "var", "1", "hayati", "kritik"}
SIRA = {"Stoksuz": 0, "Kritik altında": 1, "Sipariş yolda": 2, "Yeterli": 3, "Hesaplanamadı": 4}


def oku_stok(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, STOK, ("kod", "mevcut")):
        kod = metin(r["kod"])
        if not kod:
            continue
        d = {a: para(r[a]) if metin(r[a]) else None for a in ("mevcut", "gunluk", "aylik", "tedarik", "emniyet", "bekleyen", "fiyat")}
        if d["mevcut"] is None or any(metin(r[a]) and (d[a] is None or d[a] < 0) for a in d):
            uy.append(uyari("Yüksek", "Satır okunamadı", kod, f"Satır {r['_satir']}: sayılardan biri okunamadı ya da eksi. Malzeme alınmadı."))
            continue
        if katla(kod) in gorulen:
            uy.append(uyari("Orta", "Malzeme iki kez yazılmış", kod, f"Satır {r['_satir']}: ikinci satır alınmadı; depo bazında stok varsa toplayıp tek satır yazın."))
            continue
        gorulen.add(katla(kod))
        gunluk = d["gunluk"] if d["gunluk"] is not None else d["aylik"] / 30 if d["aylik"] is not None else None
        l.append({"kod": kod, "ad": metin(r["ad"]), "birim": metin(r["birim"]), "mevcut": d["mevcut"], "gunluk": gunluk, "tedarik": d["tedarik"], "emniyet": d["emniyet"],
                  "bekleyen": d["bekleyen"] or SIFIR, "fiyat": d["fiyat"], "hayati": katla(r["hayati"]) in EVET})
    return l, uy


def oku_tuketim(yol: Path, gun: int) -> tuple[dict, list[dict]]:
    uy, d = [], defaultdict(lambda: SIFIR)
    for r in kayitlar(yol, TUKETIM, ("kod", "miktar")):
        kod = metin(r["kod"])
        if not kod:
            continue
        m = para(r["miktar"])
        if m is None or m < 0:
            uy.append(uyari("Yüksek", "Tüketim satırı okunamadı", kod, f"Satır {r['_satir']}: miktar okunamadı. Satır alınmadı."))
            continue
        d[katla(kod)] += m
    return {k: v / gun for k, v in d.items()}, uy


def degerlendir(malzemeler: list[dict], gun: date, emniyet: Decimal = Decimal(7), kapsam: Decimal = Decimal(30), tedarik_var: Decimal | None = None, tuketim: dict | None = None) -> dict:
    if not malzemeler:
        raise ValueError("Stok kaydı yok.")
    uy, sonuc = [], []
    for m in malzemeler:
        hiz = (tuketim or {}).get(katla(m["kod"]), m["gunluk"])
        ted = m["tedarik"] if m["tedarik"] is not None else tedarik_var
        emn = m["emniyet"] if m["emniyet"] is not None else emniyet
        ad = m["ad"] or m["kod"]
        x = {**m, "hiz": yuvarla(hiz) if hiz is not None else None, "ted": ted, "emn": emn, "emniyet_stok": None, "kritik": None, "stok_gun": None, "tukenme": None, "oneri": None, "tutar": None}
        if not hiz or ted is None:
            x["durum"] = "Hesaplanamadı"
            uy.append(uyari("Bilgi", "Tüketim ya da tedarik süresi yok", m["kod"], f"{ad}: " + ("günlük tüketim" if not hiz else "tedarik süresi") + " bilinmediği için kritik seviye hesaplanmadı."))
            sonuc.append(x)
            continue
        x["emniyet_stok"] = yuvarla(hiz * emn)
        x["kritik"] = yuvarla(hiz * (ted + emn))
        x["stok_gun"] = yuvarla(m["mevcut"] / hiz)
        x["tukenme"] = gun + timedelta(days=int(m["mevcut"] / hiz))
        ihtiyac = hiz * (ted + emn + kapsam) - m["mevcut"] - m["bekleyen"]
        x["oneri"] = Decimal(math.ceil(ihtiyac)) if ihtiyac > 0 else SIFIR
        x["tutar"] = yuvarla(x["oneri"] * m["fiyat"]) if m["fiyat"] is not None else None
        if m["mevcut"] == 0:
            x["durum"] = "Stoksuz"
        elif m["mevcut"] + m["bekleyen"] < x["kritik"]:
            x["durum"] = "Kritik altında"
        elif m["mevcut"] < x["kritik"]:
            x["durum"] = "Sipariş yolda"
        else:
            x["durum"] = "Yeterli"
        if x["durum"] not in ("Stoksuz", "Kritik altında"):
            x["oneri"], x["tutar"] = SIFIR, None
        etiket = " (hayati)" if m["hayati"] else ""
        if x["durum"] == "Stoksuz":
            uy.append(uyari("Yüksek", "Stoksuz", m["kod"], f"{ad}{etiket}: stok yok; günlük tüketim {x['hiz']}." + (f" Bekleyen sipariş {m['bekleyen']:,.0f}." if m["bekleyen"] else "")))
        elif m["mevcut"] < x["emniyet_stok"]:
            uy.append(uyari("Yüksek", "Emniyet stoğunun altında", m["kod"], f"{ad}{etiket}: mevcut {m['mevcut']:,.0f}, emniyet stoğu {x['emniyet_stok']:,.0f}; tahmini tükenme {x['tukenme']:%d.%m.%Y}."))
        elif x["durum"] == "Kritik altında":
            uy.append(uyari("Yüksek" if m["hayati"] else "Orta", "Kritik seviyenin altında", m["kod"], f"{ad}{etiket}: mevcut {m['mevcut']:,.0f} + bekleyen {m['bekleyen']:,.0f} < kritik seviye {x['kritik']:,.0f}; önerilen sipariş {x['oneri']:,.0f}."))
        if x["durum"] == "Sipariş yolda" and x["stok_gun"] < ted:
            uy.append(uyari("Orta", "Sipariş gelmeden tükenebilir", m["kod"], f"{ad}{etiket}: stok {x['stok_gun']} günde biter, tedarik süresi {ted} gün. Siparişin teslim tarihini teyit edin."))
        sonuc.append(x)
    sonuc.sort(key=lambda x: (SIRA[x["durum"]], not x["hayati"], x["stok_gun"] if x["stok_gun"] is not None else Decimal(10 ** 9), katla(x["kod"])))
    ozet = {d: sum(1 for x in sonuc if x["durum"] == d) for d in SIRA}
    return {"malzemeler": sonuc, "ozet": ozet, "siparis_tutar": sum((x["tutar"] for x in sonuc if x["tutar"] and x["durum"] in ("Stoksuz", "Kritik altında")), SIFIR), "uyarilar": uy}


DURUM_RENK = {"Stoksuz": KIRMIZI, "Kritik altında": KIRMIZI, "Sipariş yolda": SARI, "Yeterli": YESIL, "Hesaplanamadı": MAVI}


def rapor_yaz(cikti: Path, r: dict, gun: date, kapsam: Decimal) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kritik Stok"
    bas = ["Durum", "Malzeme Kodu", "Malzeme Adı", "Birim", "Hayati", "Mevcut Stok", "Bekleyen Sipariş", "Günlük Tüketim", "Tedarik Süresi", "Emniyet Günü", "Emniyet Stoğu", "Kritik Seviye",
           "Stok Günü", "Tahmini Tükenme", "Önerilen Sipariş", "Sipariş Tutarı"]
    gen = (14, 13, 30, 7, 7, 9, 9, 9, 8, 8, 9, 9, 8, 11, 10, 12)

    def yaz(ws, x):
        _satir(ws, [x["durum"], x["kod"], x["ad"], x["birim"], "Evet" if x["hayati"] else "", x["mevcut"], x["bekleyen"] or None, x["hiz"], x["ted"], x["emn"], x["emniyet_stok"], x["kritik"], x["stok_gun"],
                    x["tukenme"], x["oneri"] or None, x["tutar"] or None], para_s=(8, 11, 12, 13, 16), tarih_s=(14,), dolgu={1: DURUM_RENK[x["durum"]], 5: KIRMIZI if x["hayati"] else None})
    _baslik(ws, bas, gen, "C2")
    for x in r["malzemeler"]:
        if x["durum"] in ("Stoksuz", "Kritik altında", "Sipariş yolda"):
            yaz(ws, x)
    ws.append([])
    ws.append([f"Rapor tarihi {gun:%d.%m.%Y}. Önerilen sipariş {kapsam} günlük kapsam içindir; ambalaj, minimum sipariş ve bütçe kontrolü sizindir."])
    ws = wb.create_sheet("Tüm Malzemeler")
    _baslik(ws, bas, gen, "C2")
    for x in r["malzemeler"]:
        yaz(ws, x)
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(stok_yol: Path, cikti: Path, gun: date, emniyet: Decimal = Decimal(7), kapsam: Decimal = Decimal(30), tedarik_var: Decimal | None = None, tuketim_yol: Path | None = None,
             donem: int = 90) -> dict:
    malzemeler, uy = oku_stok(stok_yol)
    tuketim, uy2 = oku_tuketim(tuketim_yol, donem) if tuketim_yol else (None, [])
    r = degerlendir(malzemeler, gun, emniyet, kapsam, tedarik_var, tuketim)
    r["uyarilar"] = sirala(uy + uy2 + r["uyarilar"])
    rapor_yaz(cikti, r, gun, kapsam)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Günlük tüketim ve tedarik süresine göre kritik stok seviyesinin altına düşen ilaç ve malzemeleri listeler.")
    p.add_argument("--stok", type=Path, help="Malzeme Kodu, Malzeme Adı, Birim, Mevcut Stok, Günlük / Aylık Tüketim, Tedarik Süresi, Emniyet Günü, Bekleyen Sipariş, Birim Fiyat, Hayati")
    p.add_argument("--tuketim", type=Path, help="Dönem tüketimi (isteğe bağlı): Malzeme Kodu, Tüketim; günlük = toplam ÷ --gun")
    p.add_argument("--gun", type=int, default=90, help="Tüketim dosyasının kapsadığı gün sayısı (varsayılan 90)")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 30.09.2026)")
    p.add_argument("--emniyet", default="7", help="Satırda yazılmamışsa emniyet günü (varsayılan 7)")
    p.add_argument("--tedarik", help="Satırda yazılmamışsa tedarik süresi, gün (varsayılan yok: hesaplanmaz)")
    p.add_argument("--kapsam", default="30", help="Önerilen siparişin karşılayacağı ek gün (varsayılan 30)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kritik_stok_raporu.xlsx")
    a = p.parse_args(argv)
    yol = a.stok or ORNEK / "eczane_stok_durumu.csv"
    gun = tarih(a.tarih) if a.tarih else (date(2026, 9, 30) if a.stok is None else date.today())
    emniyet, kapsam = para(a.emniyet), para(a.kapsam)
    tedarik_var = para(a.tedarik) if a.tedarik else None
    if gun is None or a.gun < 1 or emniyet is None or emniyet < 0 or kapsam is None or kapsam < 0 or (a.tedarik and (tedarik_var is None or tedarik_var < 0)):
        print("[X] --tarih GG.AA.YYYY, --gun en az 1, --emniyet, --tedarik ve --kapsam sıfır ya da üstü olmalı")
        return 2
    for y in (yol, a.tuketim):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(yol, a.cikti, gun, emniyet, kapsam, tedarik_var, a.tuketim, a.gun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {len(r['malzemeler'])} malzeme · stoksuz {o['Stoksuz']} · kritik altında {o['Kritik altında']} · sipariş yolda {o['Sipariş yolda']} · yeterli {o['Yeterli']}"
          + f" · önerilen sipariş tutarı {r['siparis_tutar']:,.2f}")
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
