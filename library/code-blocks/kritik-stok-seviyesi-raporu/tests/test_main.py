import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
GUN = date(2026, 9, 30)


def m(kod, mevcut, gunluk=None, tedarik=None, emniyet=None, bekleyen=0, fiyat=None, hayati=False):
    c = lambda x: None if x is None else D(x)   # noqa: E731
    return {"kod": kod, "ad": "", "birim": "", "mevcut": D(mevcut), "gunluk": c(gunluk), "tedarik": c(tedarik), "emniyet": c(emniyet), "bekleyen": D(bekleyen), "fiyat": c(fiyat), "hayati": hayati}


def veri():
    return [m("M1", 100, 10, 5, fiyat=2), m("M2", 20, 4, 10, bekleyen=60), m("M3", 0, 2, 3, hayati=True), m("M4", 100, 1, 5, 0), m("M5", 5)]


class DegerlendirTesti(unittest.TestCase):
    def test_durumlar(self):
        r = main.degerlendir(veri(), GUN)
        self.assertEqual([(x["kod"], x["durum"], x["emniyet_stok"], x["kritik"], x["stok_gun"], x["tukenme"], x["oneri"], x["tutar"]) for x in r["malzemeler"]],
                         [("M3", "Stoksuz", D("14.00"), D("20.00"), D("0.00"), GUN, D(80), None), ("M1", "Kritik altında", D("70.00"), D("120.00"), D("10.00"), date(2026, 10, 10), D(320), D("640.00")),
                          ("M2", "Sipariş yolda", D("28.00"), D("68.00"), D("5.00"), date(2026, 10, 5), D(0), None), ("M4", "Yeterli", D("0.00"), D("5.00"), D("100.00"), date(2027, 1, 8), D(0), None),
                          ("M5", "Hesaplanamadı", None, None, None, None, None, None)])
        self.assertEqual((r["ozet"], r["siparis_tutar"]), ({"Stoksuz": 1, "Kritik altında": 1, "Sipariş yolda": 1, "Yeterli": 1, "Hesaplanamadı": 1}, D("640.00")))
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Bilgi", "Tüketim ya da tedarik süresi yok", "M5"), ("Orta", "Kritik seviyenin altında", "M1"), ("Orta", "Sipariş gelmeden tükenebilir", "M2"),
                          ("Yüksek", "Emniyet stoğunun altında", "M2"), ("Yüksek", "Stoksuz", "M3")])

    def test_secenekler(self):
        r = main.degerlendir([m("M1", 100, 10, 5, fiyat=2, hayati=True), m("M6", 50, None), m("M7", 30, 1)], GUN, emniyet=D(0), kapsam=D(0), tedarik_var=D(40), tuketim={"m6": D(5)})
        self.assertEqual([(x["kod"], x["durum"], x["kritik"], x["oneri"]) for x in r["malzemeler"]],
                         [("M6", "Kritik altında", D("200.00"), D(150)), ("M7", "Kritik altında", D("40.00"), D(10)), ("M1", "Yeterli", D("50.00"), D(0))])
        with self.assertRaises(ValueError):
            main.degerlendir([], GUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "eczane_stok_durumu.csv", cikti, GUN)
            self.assertEqual((r["ozet"], r["siparis_tutar"]), ({"Stoksuz": 1, "Kritik altında": 1, "Sipariş yolda": 2, "Yeterli": 7, "Hesaplanamadı": 1}, D("1286.40")))
            self.assertEqual([(x["kod"], x["oneri"]) for x in r["malzemeler"] if x["oneri"]], [("IL-0710", D(32)), ("IL-0520", D(32))])
            self.assertEqual([(u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Emniyet stoğunun altında", "TM-2020"), ("Stoksuz", "IL-0710"), ("Kritik seviyenin altında", "IL-0520"), ("Sipariş gelmeden tükenebilir", "IL-0330"),
                              ("Tüketim ya da tedarik süresi yok", "TM-2010")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Kritik Stok", "Tüm Malzemeler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, z = Path(tmp) / "s.csv", Path(tmp) / "t.csv"
            y.write_text("İlaç Kodu;Eldeki Stok;Aylık Tüketim;Temin Süresi;Açık Sipariş;Hayati\nA;1.200;300;5;100;Evet\na;5;;;;\nB;x;;;;\nC;5;-1;;;\nD;5;;;;\n;5;;;;\n", encoding="utf-8")
            z.write_text("İlaç Kodu;Tüketim\nD;900\nd;90\nE;x\n", encoding="utf-8")
            l, uy = main.oku_stok(y)
            self.assertEqual(([(x["kod"], x["mevcut"], x["gunluk"], x["tedarik"], x["bekleyen"], x["hayati"]) for x in l], [u["tur"] for u in uy]),
                             ([("A", D(1200), D(10), D(5), D(100), True), ("D", D(5), None, None, D(0), False)], ["Malzeme iki kez yazılmış", "Satır okunamadı", "Satır okunamadı"]))
            t, uy = main.oku_tuketim(z, 90)
            self.assertEqual((t, len(uy)), ({"d": D(11)}, 1))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--stok", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tuketim", str(z), "--gun", "90", "--tarih", "30.09.2026", "--emniyet", "5", "--tedarik", "7", "--kapsam", "15"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--gun", "0"], ["--emniyet", "-1"], ["--tedarik", "x"])], [2] * 4)
            self.assertEqual(main.main(arg + ["--tuketim", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("İlaç Kodu;Hayati\nA;x\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
