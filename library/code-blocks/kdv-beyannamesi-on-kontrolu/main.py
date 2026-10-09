"""
KDV Beyannamesi Ön Kontrolü — Workers / Workless kod bloğu
Muhasebe › Genel Muhasebe Uzmanı

Satış ve alış fatura listelerini, mizanı ve hazırlanan KDV beyannamesi tutarlarını karşılaştırır; uyumsuzlukları
beyan öncesi yakalar. Beyanname alanları sade bir anahtar–tutar listesiyle verilir (bkz. README).
  - Liste kontrolleri: dönem dışı tarih, mükerrer fatura, KDV ≠ matrah × oran, oran tarihte geçerli mi,
    tevkif edilen KDV = KDV × tevkifat oranı, indirim süresi (KDVK md. 29/3: vergiyi doğuran olayın ait olduğu
    takvim yılını izleyen yıl sonu) geçmiş alış faturası, indirilemez işaretli alışlar.
  - Satış tarafı: oran bazında matrah ve KDV (tevkifatsız) ↔ beyanname; kısmi tevkifatlı işlemler (satıcının beyan
    ettiği KDV = KDV − tevkif edilen) ↔ beyanname; istisna matrahı ↔ beyanname; listenin beyan edilecek KDV'si ↔ 391.
  - Alış tarafı: indirilebilir ve süresi geçmemiş alış KDV'si ↔ 191 ↔ beyanname "bu döneme ait indirilecek KDV".
  - Devreden: beyanname "önceki dönemden devreden" ↔ 190.
  - Beyanname aritmetiği: toplam hesaplanan = oran KDV'leri + kısmi tevkifat KDV'si; ödenecek / sonraki döneme
    devreden = hesaplanan − (devreden + indirilecek).
  - Gelir mutabakatı: 600 + 601 + 602 − 610 − 611 − 612 ↔ satış listesi matrahı (bilgi).
Rapor: mutabakat tablosu, beyanname aritmetiği, liste bulguları, oran özetleri, sorumlu sıfatıyla KDV, uyarılar.
İnternete bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: Eylül 2026
    python main.py --satislar s.xlsx --alislar a.xlsx --mizan mizan.xlsx --beyanname beyan.csv --donem 2026-09
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
K2 = Decimal("0.01")
SIFIR = Decimal(0)
ORAN_DEGISIMI = date(2023, 7, 10)
ORANLAR_YENI = {Decimal(0), Decimal(1), Decimal(10), Decimal(20)}
ORANLAR_ESKI = {Decimal(0), Decimal(1), Decimal(8), Decimal(18)}

LISTE_SUTUNLARI = {"no": ("fatura no", "belge no"), "tarih": ("tarih", "fatura tarihi"), "taraf": ("alici", "satici", "unvan", "cari"),
                   "matrah": ("matrah", "tutar", "kdv matrahi"), "oran": ("kdv orani",), "kdv": ("kdv tutari", "kdv"),
                   "t_oran": ("tevkifat orani",), "t_tutar": ("tevkif edilen kdv", "sorumlu sifatiyla beyan edilen kdv", "tevkifat tutari"),
                   "tur": ("islem turu", "tur"), "indirilebilir": ("indirilebilir", "kdv indirilebilir")}
MIZAN_SUTUNLARI = {"kod": ("hesap kodu", "hesap", "kod"), "ad": ("hesap adi",), "borc": ("borc", "borc toplami"), "alacak": ("alacak", "alacak toplami")}
BEYAN_ANAHTAR = {"toplam hesaplanan kdv": "hesaplanan", "onceki donemden devreden kdv": "devreden", "onceki donemden devreden indirilecek kdv": "devreden",
                 "bu doneme ait indirilecek kdv": "indirilecek", "odenecek kdv": "odenecek", "odenmesi gereken kdv": "odenecek",
                 "sonraki doneme devreden kdv": "sonraki", "kismi tevkifat matrahi": "t_matrah", "kismi tevkifat beyan edilen kdv": "t_kdv",
                 "istisna matrahi": "istisna"}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("%", "").replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tevkifat_orani(x) -> Decimal | None:
    if x in (None, ""):
        return None
    m = re.fullmatch(r"\s*(\d+)\s*/\s*(\d+)\s*", str(x))
    if m:
        return Decimal(m[1]) / Decimal(m[2]) if int(m[2]) else None
    v = para(x)
    return None if v is None else v / 100 if v > 1 else v


def yuv(x: Decimal) -> Decimal:
    return x.quantize(K2, ROUND_HALF_UP)


def tl(x: Decimal | None) -> str:
    return "—" if x is None else f"{yuv(x):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Fatura:
    taraf: str                  # satis | alis
    no: str
    tarih: date | None
    karsi: str
    matrah: Decimal
    oran: Decimal
    kdv: Decimal
    t_oran: Decimal | None
    t_tutar: Decimal
    istisna: bool
    indirilebilir: bool
    satir: int
    bulgular: list = field(default_factory=list)
    hesaba_dahil: bool = True

    @property
    def beyan_kdv(self) -> Decimal:
        """Satıcının beyan ettiği KDV (kısmi tevkifatta tevkif edilen kısım hariç)."""
        return self.kdv - self.t_tutar


def oku_liste(yol: Path, taraf: str) -> tuple[list[Fatura], list[dict]]:
    sonuc, uy = [], []
    for r in kayitlar(yol, LISTE_SUTUNLARI, ("no", "matrah", "kdv")):
        m, k = para(r.get("matrah")), para(r.get("kdv"))
        if not metin(r.get("no")) or m is None or k is None:
            if any(r.get(x) for x in ("no", "matrah")):
                uy.append({"onem": "Orta", "tur": "Okunamayan satır", "kaynak": yol.name, "aciklama": f"Satır {r['_satir']}: matrah veya KDV okunamadı"})
            continue
        oran = para(r.get("oran"))
        if oran is None:
            oran = yuv(k / m * 100) if m else SIFIR
        tur = katla(r.get("tur"))
        sonuc.append(Fatura(taraf, metin(r["no"]), tarih(r.get("tarih")), metin(r.get("taraf")), m, oran, k, tevkifat_orani(r.get("t_oran")),
                            para(r.get("t_tutar")) or SIFIR, tur.startswith(("istisna", "ihracat")),
                            not katla(r.get("indirilebilir")).startswith(("h", "0")), r["_satir"]))
    return sonuc, uy


def oku_mizan(yol: Path | None) -> dict[str, tuple[Decimal, Decimal]]:
    if not yol:
        return {}
    m = {}
    for r in kayitlar(yol, MIZAN_SUTUNLARI, ("kod", "borc", "alacak")):
        k = metin(r.get("kod"))
        if k:
            b0, a0 = m.get(k, (SIFIR, SIFIR))
            m[k] = (b0 + (para(r.get("borc")) or SIFIR), a0 + (para(r.get("alacak")) or SIFIR))
    return m


def mizan_bakiye(m: dict, onekler: tuple[str, ...], yon: str) -> Decimal | None:
    """Ana hesap öneklerine göre bakiye; alt hesaplı mizanda yalnız yaprak hesaplar toplanır."""
    kodlar = [k for k in m if k.split(".")[0] in onekler]
    if not kodlar:
        return None
    yaprak = [k for k in kodlar if not any(o != k and o.startswith(k + ".") for o in kodlar)]
    b = sum((m[k][0] for k in yaprak), SIFIR)
    a = sum((m[k][1] for k in yaprak), SIFIR)
    return b - a if yon == "borc" else a - b


def oku_beyanname(yol: Path | None) -> tuple[dict, dict]:
    if not yol:
        return {}, {}
    alanlar, oranlar = {}, defaultdict(dict)
    for r in tablo_oku(yol)[1:]:
        if len(r) < 2:
            continue
        ad, v = katla(r[0]), para(r[1])
        if v is None:
            continue
        mt = re.fullmatch(r"(matrah|kdv) (\d+)", ad)
        if mt:
            oranlar[Decimal(mt[2])][mt[1]] = v
        elif ad in BEYAN_ANAHTAR:
            alanlar[BEYAN_ANAHTAR[ad]] = v
    return alanlar, dict(oranlar)


# ----------------------------------------------------------------------------
# Kontroller
# ----------------------------------------------------------------------------

def liste_kontrol(faturalar: list[Fatura], bas: date, bit: date, tol: Decimal) -> None:
    sayac = Counter((f.taraf, f.no, katla(f.karsi)) for f in faturalar)
    goruldu = set()
    for f in faturalar:
        b = f.bulgular
        anahtar = (f.taraf, f.no, katla(f.karsi))
        if sayac[anahtar] > 1:
            if anahtar in goruldu:
                b.append(("Yüksek", "Mükerrer fatura", f"{f.no} listede {sayac[anahtar]} kez var; tekrarlar hesaplamaya dahil edildi"))
            goruldu.add(anahtar)
        if f.tarih and not (bas <= f.tarih <= bit) and not (f.taraf == "alis" and f.tarih < bas):
            b.append(("Yüksek", "Dönem dışı", f"Fatura tarihi {f.tarih:%d.%m.%Y} beyan dönemi dışında; hesaplamaya dahil edilmedi"))
            f.hesaba_dahil = False
        gecerli = ORANLAR_YENI if not f.tarih or f.tarih >= ORAN_DEGISIMI else ORANLAR_ESKI
        if f.oran not in gecerli:
            b.append(("Yüksek", "Geçersiz oran", f"%{f.oran:g} fatura tarihinde geçerli bir KDV oranı değil"))
        beklenen = yuv(f.matrah * f.oran / 100)
        if abs(beklenen - f.kdv) > tol:
            b.append(("Yüksek", "KDV hesabı", f"{tl(f.matrah)} × %{f.oran:g} = {tl(beklenen)}; listede {tl(f.kdv)} (fark {tl(f.kdv - beklenen)})"))
        if f.t_oran is not None:
            bt = yuv(f.kdv * f.t_oran)
            if abs(bt - f.t_tutar) > tol:
                b.append(("Yüksek", "Tevkifat hesabı", f"{tl(f.kdv)} × {f.t_oran * 10:g}/10 = {tl(bt)}; listede {tl(f.t_tutar)}"))
        if f.istisna and f.kdv:
            b.append(("Yüksek", "İstisnada KDV", f"İstisna işaretli işlemde KDV {tl(f.kdv)}"))
        if f.taraf == "alis":
            if f.tarih and f.tarih.year < bas.year - 1:
                b.append(("Yüksek", "İndirim süresi", f"{f.tarih:%d.%m.%Y} tarihli faturanın KDV'si için indirim hakkı {f.tarih.year + 1} yılı sonunda "
                          "sona ermiş olabilir (KDVK md. 29/3); indirilecek KDV'ye dahil edilmedi"))
                f.hesaba_dahil = False
            elif f.tarih and f.tarih < bas:
                b.append(("Bilgi", "Önceki dönem faturası", f"{f.tarih:%d.%m.%Y} tarihli fatura bu dönemde indiriliyor; indirim süresi içinde"))
            if not f.indirilebilir:
                b.append(("Bilgi", "İndirilemez KDV", f"{tl(f.kdv)} KDV indirilemez işaretli; indirilecek KDV'ye dahil edilmedi"))


@dataclass
class Satir:
    grup: str
    kalem: str
    liste: Decimal | None
    mizan: Decimal | None
    beyan: Decimal | None
    not_: str = ""

    def fark(self, a, b):
        return None if a is None or b is None else a - b


def mutabakat(faturalar: list[Fatura], m: dict, alanlar: dict, oranlar: dict, tol: Decimal) -> tuple[list[Satir], list[dict], dict]:
    s = [f for f in faturalar if f.taraf == "satis" and f.hesaba_dahil]
    a = [f for f in faturalar if f.taraf == "alis" and f.hesaba_dahil]
    tablo, uy = [], []
    normal = [f for f in s if not f.istisna and f.t_oran is None]
    tum_oran = sorted({f.oran for f in normal if f.oran} | {o for o in oranlar if o})
    for o in tum_oran:
        lst = [f for f in normal if f.oran == o]
        tablo.append(Satir("Satış", f"Matrah %{o:g}", sum((f.matrah for f in lst), SIFIR), None, oranlar.get(o, {}).get("matrah")))
        tablo.append(Satir("Satış", f"KDV %{o:g}", sum((f.kdv for f in lst), SIFIR), None, oranlar.get(o, {}).get("kdv")))
    tev = [f for f in s if f.t_oran is not None]
    tablo.append(Satir("Satış", "Kısmi tevkifat matrahı", sum((f.matrah for f in tev), SIFIR), None, alanlar.get("t_matrah")))
    tablo.append(Satir("Satış", "Kısmi tevkifat — satıcının beyan ettiği KDV", sum((f.beyan_kdv for f in tev), SIFIR), None, alanlar.get("t_kdv")))
    tablo.append(Satir("Satış", "İstisna matrahı", sum((f.matrah for f in s if f.istisna), SIFIR), None, alanlar.get("istisna")))
    hesaplanan = sum((f.beyan_kdv for f in s), SIFIR)
    tablo.append(Satir("Satış", "Toplam hesaplanan KDV (391)", hesaplanan, mizan_bakiye(m, ("391",), "alacak"), alanlar.get("hesaplanan")))
    indirilecek = sum((f.kdv for f in a if f.indirilebilir), SIFIR)
    tablo.append(Satir("Alış", "Bu döneme ait indirilecek KDV (191)", indirilecek, mizan_bakiye(m, ("191",), "borc"), alanlar.get("indirilecek")))
    tablo.append(Satir("Alış", "Önceki dönemden devreden KDV (190)", None, mizan_bakiye(m, ("190",), "borc"), alanlar.get("devreden")))
    gelir = None
    if m:
        g1 = mizan_bakiye(m, ("600", "601", "602"), "alacak")
        g2 = mizan_bakiye(m, ("610", "611", "612"), "borc")
        if g1 is not None:
            gelir = g1 - (g2 or SIFIR)
    tablo.append(Satir("Gelir", "Net satışlar (600–602 − 610–612) ↔ satış matrahı", sum((f.matrah for f in s), SIFIR), gelir, None,
                       "Bilgi: KDV'siz gelirler, iadeler ve dönemsellik farkı oluşturabilir"))
    for t in tablo:
        degerler = [x for x in (t.liste, t.mizan, t.beyan) if x is not None]
        if len(degerler) >= 2 and max(degerler) - min(degerler) > tol:
            onem = "Bilgi" if t.grup == "Gelir" else "Yüksek"
            parcalar = [f"{ad} {tl(v)}" for ad, v in (("liste", t.liste), ("mizan", t.mizan), ("beyanname", t.beyan)) if v is not None]
            uy.append({"onem": onem, "tur": "Uyumsuzluk", "kaynak": t.kalem, "aciklama": ", ".join(parcalar) + f" (en büyük fark {tl(max(degerler) - min(degerler))})"})
    toplamlar = {"hesaplanan": hesaplanan, "indirilecek": indirilecek, "sorumlu": sum((f.t_tutar for f in a if f.t_oran is not None), SIFIR)}
    return tablo, uy, toplamlar


def aritmetik(alanlar: dict, oranlar: dict, tol: Decimal) -> tuple[list[tuple], list[dict]]:
    uy, sat = [], []
    if not alanlar:
        return sat, uy
    h = sum((v.get("kdv", SIFIR) for v in oranlar.values()), SIFIR) + alanlar.get("t_kdv", SIFIR)
    sat.append(("Oran KDV'leri + kısmi tevkifat KDV'si", h, alanlar.get("hesaplanan")))
    for o, v in sorted(oranlar.items()):
        if "matrah" in v and "kdv" in v and abs(yuv(v["matrah"] * o / 100) - v["kdv"]) > tol:
            uy.append({"onem": "Yüksek", "tur": "Beyanname aritmetiği", "kaynak": f"%{o:g}", "aciklama": f"Matrah {tl(v['matrah'])} × %{o:g} = "
                       f"{tl(yuv(v['matrah'] * o / 100))}; beyannamede KDV {tl(v['kdv'])}"})
    hes = alanlar.get("hesaplanan", h)
    indirim = alanlar.get("devreden", SIFIR) + alanlar.get("indirilecek", SIFIR)
    sat.append(("Ödenecek KDV", max(hes - indirim, SIFIR), alanlar.get("odenecek")))
    sat.append(("Sonraki döneme devreden KDV", max(indirim - hes, SIFIR), alanlar.get("sonraki")))
    for ad, hesap, beyan in sat:
        if beyan is not None and abs(hesap - beyan) > tol:
            uy.append({"onem": "Yüksek", "tur": "Beyanname aritmetiği", "kaynak": ad, "aciklama": f"Hesaplanan {tl(hesap)}; beyannamede {tl(beyan)}"})
    return sat, uy


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, YESIL = PatternFill("solid", fgColor="FDE2E1"), PatternFill("solid", fgColor="E3F4E1")
UST = Alignment(vertical="top", wrap_text=True)
PF = "#,##0.00"


def _baslik(ws, basliklar, genislik=()):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def _f(x):
    return None if x is None else float(x)


def rapor_yaz(cikti: Path, s: dict) -> None:
    tol = s["tol"]
    wb = Workbook()
    mt = wb.active
    mt.title = "Mutabakat"
    _baslik(mt, ["Grup", "Kalem", "Listeler", "Mizan", "Beyanname", "Liste − Beyanname", "Mizan − Beyanname", "Liste − Mizan", "Durum", "Not"],
            (8, 44, 15, 15, 15, 15, 15, 15, 9, 40))
    for t in s["tablo"]:
        f1, f2, f3 = t.fark(t.liste, t.beyan), t.fark(t.mizan, t.beyan), t.fark(t.liste, t.mizan)
        farklar = [abs(x) for x in (f1, f2, f3) if x is not None]
        durum = "—" if not farklar else "Tamam" if max(farklar) <= tol else "Fark"
        mt.append([t.grup, t.kalem, _f(t.liste), _f(t.mizan), _f(t.beyan), _f(f1), _f(f2), _f(f3), durum, t.not_])
        r = mt.max_row
        for j in range(3, 9):
            mt.cell(r, j).number_format = PF
        if durum != "—":
            mt.cell(r, 9).fill = YESIL if durum == "Tamam" else (KIRMIZI if t.grup != "Gelir" else PatternFill("solid", fgColor="E8F0FE"))

    ar = wb.create_sheet("Beyanname Aritmetiği")
    _baslik(ar, ["Kalem", "Hesaplanan", "Beyannamede", "Fark"], (40, 16, 16, 14))
    for ad, h, b in s["aritmetik"]:
        ar.append([ad, float(h), _f(b), _f(None if b is None else b - h)])
        for j in (2, 3, 4):
            ar.cell(ar.max_row, j).number_format = PF
    ar.append([])
    ar.append(["Not: Sonuç yalnız bu sade alan listesiyle hesaplanır. Beyannamedeki diğer indirimler, iade ve mahsup talepleri, ihraç kayıtlı "
               "teslimler ve özel matrah satırları ayrıca kontrol edilmelidir."])

    bl = wb.create_sheet("Liste Bulguları")
    _baslik(bl, ["Önem", "Tür", "Liste", "Fatura No", "Tarih", "Karşı Taraf", "Matrah", "KDV", "Açıklama", "Düzeltme"], (8, 20, 7, 19, 11, 22, 13, 12, 70, 16))
    for f in s["faturalar"]:
        for o, t, a in f.bulgular:
            bl.append([o, t, "Satış" if f.taraf == "satis" else "Alış", f.no, f.tarih, f.karsi, float(f.matrah), float(f.kdv), a, ""])
            r = bl.max_row
            bl.cell(r, 1).fill = PatternFill("solid", fgColor=RENK[o])
            bl.cell(r, 5).number_format = "DD.MM.YYYY"
            for j in (7, 8):
                bl.cell(r, j).number_format = PF
            bl.cell(r, 9).alignment = UST
            bl.cell(r, 10).fill = PatternFill("solid", fgColor="FFF4CE")

    oz = wb.create_sheet("Oran Özeti")
    _baslik(oz, ["Liste", "Oran", "Fatura", "Matrah", "KDV", "Tevkif Edilen", "Hesaba Dahil KDV"], (8, 7, 8, 16, 14, 14, 16))
    for taraf, ad in (("satis", "Satış"), ("alis", "Alış")):
        g = defaultdict(list)
        for f in s["faturalar"]:
            if f.taraf == taraf:
                g[f.oran].append(f)
        for o, lst in sorted(g.items()):
            dahil = [f for f in lst if f.hesaba_dahil and (taraf == "satis" or f.indirilebilir)]
            oz.append([ad, f"%{o:g}", len(lst), float(sum(f.matrah for f in lst)), float(sum(f.kdv for f in lst)), float(sum(f.t_tutar for f in lst)),
                       float(sum((f.beyan_kdv if taraf == "satis" else f.kdv) for f in dahil))])
            for j in (4, 5, 6, 7):
                oz.cell(oz.max_row, j).number_format = PF
    oz.append([])
    oz.append([f"Sorumlu sıfatıyla beyan edilecek KDV (alışlardaki tevkifat, 2 No'lu KDV beyannamesi): {tl(s['toplamlar']['sorumlu'])} TL"])

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Kaynak", "Açıklama"], (9, 22, 34, 90))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["kaynak"], u["aciklama"]])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def donem_ayir(x: str) -> tuple[date, date] | None:
    m = re.fullmatch(r"\s*(\d{4})[-./](\d{1,2})\s*|\s*(\d{1,2})[-./](\d{4})\s*", x or "")
    if not m:
        return None
    y, a = (int(m[1]), int(m[2])) if m[1] else (int(m[4]), int(m[3]))
    if not 1 <= a <= 12:
        return None
    son = date(y + (a == 12), a % 12 + 1, 1)
    return date(y, a, 1), date.fromordinal(son.toordinal() - 1)


def calistir(satis_yolu: Path, alis_yolu: Path, cikti: Path, donem: tuple[date, date], mizan_yolu: Path | None = None,
             beyan_yolu: Path | None = None, tol: Decimal = Decimal("0.05")) -> dict:
    satislar, u1 = oku_liste(satis_yolu, "satis")
    alislar, u2 = oku_liste(alis_yolu, "alis")
    faturalar = satislar + alislar
    if not faturalar:
        raise ValueError("Satış ve alış listelerinde fatura bulunamadı")
    liste_kontrol(faturalar, *donem, tol)
    m = oku_mizan(mizan_yolu)
    alanlar, oranlar = oku_beyanname(beyan_yolu)
    tablo, u3, toplamlar = mutabakat(faturalar, m, alanlar, oranlar, tol)
    arit, u4 = aritmetik(alanlar, oranlar, tol)
    uyarilar = u1 + u2 + u3 + u4 + [{"onem": o, "tur": t, "kaynak": f.no, "aciklama": a} for f in faturalar for o, t, a in f.bulgular]
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    uyarilar.sort(key=lambda u: (sira[u["onem"]], u["tur"] != "Uyumsuzluk", u["tur"], u["kaynak"]))
    s = {"faturalar": faturalar, "tablo": tablo, "aritmetik": arit, "uyarilar": uyarilar, "toplamlar": toplamlar, "tol": tol, "donem": donem}
    rapor_yaz(cikti, s)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Satış / alış listeleri, mizan ve KDV beyannamesi tutarlarını karşılaştırır, beyan öncesi uyumsuzlukları bulur.")
    p.add_argument("--satislar", type=Path, default=ORNEK / "satislar.csv",
                   help="Fatura No, Tarih, Alıcı, Matrah, KDV Oranı, KDV Tutarı, Tevkifat Oranı, Tevkif Edilen KDV, İşlem Türü (İstisna)")
    p.add_argument("--alislar", type=Path, default=ORNEK / "alislar.csv",
                   help="Fatura No, Tarih, Satıcı, Matrah, KDV Oranı, KDV Tutarı, İndirilebilir, Tevkifat Oranı, Sorumlu Sıfatıyla Beyan Edilen KDV")
    p.add_argument("--mizan", type=Path, help="İsteğe bağlı: Hesap Kodu, Hesap Adı, Borç, Alacak (dönem)")
    p.add_argument("--beyanname", type=Path, help="İsteğe bağlı: Alan;Tutar listesi (bkz. README)")
    p.add_argument("--donem", help="Beyan dönemi YYYY-AA (örnek veride 2026-09)")
    p.add_argument("--tolerans", type=float, default=0.05, help="Tutarlar arası kabul edilen fark, TL (varsayılan 0,05)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kdv_beyanname_on_kontrol.xlsx")
    a = p.parse_args(argv)
    ornek = a.satislar == ORNEK / "satislar.csv"
    miz = a.mizan or (ORNEK / "mizan.csv" if ornek else None)
    bey = a.beyanname or (ORNEK / "beyanname.csv" if ornek else None)
    for y in (a.satislar, a.alislar, miz, bey):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    donem = donem_ayir(a.donem) if a.donem else (donem_ayir("2026-09") if ornek else None)
    if not donem:
        print("[X] --donem YYYY-AA biçiminde olmalı (ör. 2026-09)")
        return 2
    try:
        s = calistir(a.satislar, a.alislar, a.cikti, donem, miz, bey, Decimal(str(a.tolerans)))
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = s["toplamlar"]
    print(f"[OK] {donem[0]:%m.%Y} · listelerden hesaplanan KDV {tl(t['hesaplanan'])} TL · indirilecek KDV {tl(t['indirilecek'])} TL · "
          f"sorumlu sıfatıyla {tl(t['sorumlu'])} TL")
    for u in s["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['kaynak']} · {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
