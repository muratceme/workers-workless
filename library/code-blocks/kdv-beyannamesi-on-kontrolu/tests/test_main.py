import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
EYLUL = (date(2026, 9, 1), date(2026, 9, 30))


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "k.xlsx"
        cls.s = main.calistir(ORNEK / "satislar.csv", ORNEK / "alislar.csv", cls.cikti, EYLUL, ORNEK / "mizan.csv", ORNEK / "beyanname.csv")
        cls.t = {t.kalem: t for t in cls.s["tablo"]}
        cls.b = {(u["kaynak"], u["tur"]) for u in cls.s["uyarilar"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_liste_bulgulari(self):
        self.assertTrue({("SAT2026000000909", "Dönem dışı"), ("SAT2026000000908", "Mükerrer fatura"), ("TED2026000004416", "KDV hesabı"),
                         ("TED2024000001207", "İndirim süresi"), ("TED2026000004415", "İndirilemez KDV")} <= self.b)
        self.assertEqual(sum(1 for u in self.s["uyarilar"] if u["tur"] == "Mükerrer fatura"), 1)

    def test_mutabakat(self):
        t = self.t["Toplam hesaplanan KDV (391)"]
        self.assertEqual((t.liste, t.mizan, t.beyan), (D("165820.00"), D("159420.00"), D("159420.00")))   # fark = mükerrer 6.400
        t = self.t["Bu döneme ait indirilecek KDV (191)"]
        self.assertEqual((t.liste, t.mizan, t.beyan), (D("86255.00"), D("86055.00"), D("92255.00")))
        self.assertEqual(self.t["Kısmi tevkifat — satıcının beyan ettiği KDV"].liste, D("30000.00"))
        self.assertEqual(self.t["Matrah %20"].liste, D("565000"))                    # tevkifatlı ve istisna hariç
        self.assertIn(("Önceki dönemden devreden KDV (190)", "Uyumsuzluk"), self.b)
        self.assertNotIn(("Matrah %20", "Uyumsuzluk"), self.b)
        self.assertEqual(self.s["toplamlar"]["sorumlu"], D("7200.00"))

    def test_aritmetik_ve_excel(self):
        self.assertEqual([x[1] for x in self.s["aritmetik"]], [D("159420.00"), D("41765.00"), D("0")])
        self.assertFalse(any(u["tur"] == "Beyanname aritmetiği" for u in self.s["uyarilar"]))
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Mutabakat", "Beyanname Aritmetiği", "Liste Bulguları", "Oran Özeti", "Uyarılar"])


class KuralTesti(unittest.TestCase):
    def test_donem_ve_aritmetik_hatasi(self):
        self.assertEqual(main.donem_ayir("2026-02"), (date(2026, 2, 1), date(2026, 2, 28)))
        self.assertEqual(main.donem_ayir("12.2026"), (date(2026, 12, 1), date(2026, 12, 31)))
        self.assertIsNone(main.donem_ayir("2026-13"))
        sat, uy = main.aritmetik({"hesaplanan": D("100"), "devreden": D("30"), "indirilecek": D("90"), "odenecek": D("0"), "sonraki": D("10")},
                                 {D(20): {"matrah": D("500"), "kdv": D("90")}}, D("0.05"))
        turler = [u["kaynak"] for u in uy]
        self.assertIn("%20", turler)                                                 # 500 × %20 = 100 ≠ 90
        self.assertIn("Sonraki döneme devreden KDV", turler)                          # 120 − 100 = 20 ≠ 10

    def test_mizan_alt_hesap(self):
        m = {"191": (D("100"), D("0")), "191.01": (D("60"), D("0")), "191.02": (D("40"), D("0")), "600": (D("0"), D("10"))}
        self.assertEqual(main.mizan_bakiye(m, ("191",), "borc"), D("100"))
        self.assertIsNone(main.mizan_bakiye(m, ("190",), "borc"))

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--donem", "eylül"]), 2)
            self.assertEqual(main.main(["--alislar", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
