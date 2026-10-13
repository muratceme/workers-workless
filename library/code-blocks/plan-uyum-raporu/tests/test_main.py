import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)


def e(no, hat, urun, gercek, pb, gb=None, neden="", plan=100):
    g = lambda n: None if n is None else date(2026, 10, n)   # noqa: E731
    return {"no": no, "urun": urun, "hat": hat, "plan": D(plan), "gercek": D(gercek), "p_bitis": g(pb), "g_bitis": g(gb), "neden": neden}


def veri():
    return [e("E1", "H1", "Ü1", 100, 5, 5), e("E2", "H1", "Ü1", 80, 6, 6, "Arıza"), e("E3", "h1", "Ü2", 100, 6, 8, "Malzeme"), e("E4", "H2", "Ü2", 70, 7, 9, "ARIZA"), e("E5", "H2", "Ü3", 120, 8, 8), e("E6", "H2", "Ü3", 40, 10),
            e("E7", "H2", "Ü3", 0, 15)]


class UyumTesti(unittest.TestCase):
    def test_durumlar(self):
        r = main.uyum(veri(), BUGUN)
        self.assertEqual({x["no"]: (x["durum"], x["fark"], x["oran"], x["gecikme"], x["eksik"]) for x in r["emirler"]},
                         {"E1": ("Plana uygun", D(0), D("100.00"), 0, D(0)), "E2": ("Eksik", D(-20), D("80.00"), 0, D(20)), "E3": ("Geç", D(0), D("100.00"), 2, D(0)), "E4": ("Eksik ve geç", D(-30), D("70.00"), 2, D(30)),
                          "E5": ("Fazla üretim", D(20), D("120.00"), 0, D(0)), "E6": ("Gecikmiş (açık)", D(-60), D("40.00"), 2, D(60)), "E7": ("Devam ediyor", D(-100), D("0.00"), 0, D(100))})
        self.assertEqual([x["no"] for x in r["emirler"]], ["E6", "E4", "E2", "E3", "E5", "E1", "E7"])
        self.assertEqual(r["ozet"], {"emir": 6, "devam": 1, "uygun": 1, "uyum": D("16.67"), "plan": D(600), "gercek": D(510), "miktar_uyum": D("81.67"), "eksik": D(110), "fazla": D(20), "gec": 3, "zamaninda": D("50.00")})
        self.assertEqual([(s["ad"], s["emir"], s["devam"], s["uygun"], s["uyum"], s["plan"], s["gercek"], s["miktar_uyum"], s["eksik"], s["gec"], s["zamaninda"]) for s in r["hatlar"]],
                         [("H2", 3, 1, 0, D("0.00"), D(300), D(230), D("70.00"), D(90), 2, D("33.33")), ("H1", 3, 0, 1, D("33.33"), D(300), D(280), D("93.33"), D(20), 1, D("66.67"))])
        self.assertEqual([(s["ad"], s["emir"], s["miktar_uyum"]) for s in r["urunler"]], [("Ü3", 2, D("70.00")), ("Ü2", 2, D("85.00")), ("Ü1", 2, D("90.00"))])
        self.assertEqual([(n["ad"], n["emir"], n["pay"], n["eksik"], n["gecikme"], n["emirler"]) for n in r["nedenler"]],
                         [("Arıza", 2, D("40.00"), D(50), 2, ["E2", "E4"]), ("Malzeme", 1, D("20.00"), D(0), 2, ["E3"]), (main.YOK, 2, D("40.00"), D(60), 2, ["E5", "E6"])])
        self.assertEqual(([(d["durum"], d["emir"]) for d in r["durumlar"]], sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"])),
                         ([("Gecikmiş (açık)", 1), ("Eksik ve geç", 1), ("Eksik", 1), ("Geç", 1), ("Fazla üretim", 1), ("Plana uygun", 1), ("Devam ediyor", 1)], [("Orta", "Sapma nedeni yazılmamış", main.YOK), ("Yüksek", "Gecikmiş açık iş emri", "E6")]))

    def test_tolerans_ve_sinirlar(self):
        d = lambda **kw: {x["no"]: x["durum"] for x in main.uyum(veri(), BUGUN, **kw)["emirler"]}   # noqa: E731
        self.assertEqual((d(tolerans=D(20)), main.uyum(veri(), BUGUN, tolerans=D(20))["ozet"]["uyum"]),
                         ({"E1": "Plana uygun", "E2": "Plana uygun", "E3": "Geç", "E4": "Eksik ve geç", "E5": "Plana uygun", "E6": "Gecikmiş (açık)", "E7": "Devam ediyor"}, D("50.00")))
        self.assertEqual((d(tolerans=D("19.9"))["E2"], d(tolerans=D("19.9"))["E5"]), ("Eksik", "Fazla üretim"))
        s = lambda pb, gb=None: main.uyum([e("a", "H", "Ü", 100, pb, gb)], BUGUN)["emirler"][0]["durum"]   # noqa: E731
        self.assertEqual((s(12), s(11), s(12, 11), s(12, 13)), ("Devam ediyor", "Gecikmiş (açık)", "Plana uygun", "Geç"))
        o = main.uyum([e("a", "H", "Ü", 0, 20)], BUGUN)
        self.assertEqual((o["ozet"], o["nedenler"], o["uyarilar"]), ({"emir": 0, "devam": 1, "uygun": 0, "uyum": None, "plan": D(0), "gercek": D(0), "miktar_uyum": None, "eksik": D(0), "fazla": D(0), "gec": 0, "zamaninda": None}, [], []))
        with self.assertRaises(ValueError):
            main.uyum([], BUGUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "is_emirleri.csv", cikti, BUGUN)
            self.assertEqual([(x["no"], x["durum"]) for x in r["emirler"]],
                             [("IE-1008", "Gecikmiş (açık)"), ("IE-1005", "Eksik ve geç"), ("IE-1002", "Eksik"), ("IE-1007", "Eksik"), ("IE-1003", "Geç"), ("IE-1006", "Fazla üretim"), ("IE-1001", "Plana uygun"), ("IE-1004", "Plana uygun"),
                              ("IE-1010", "Plana uygun"), ("IE-1009", "Devam ediyor")])
            self.assertEqual(r["ozet"], {"emir": 9, "devam": 1, "uygun": 3, "uyum": D("33.33"), "plan": D(26200), "gercek": D(25150), "miktar_uyum": D("94.85"), "eksik": D(1350), "fazla": D(300), "gec": 3, "zamaninda": D("66.67")})
            self.assertEqual([(s["ad"], s["emir"], s["uygun"], s["plan"], s["miktar_uyum"]) for s in r["hatlar"]], [("CNC 1", 3, 1, D(2200), D("70.45")), ("Pres 1", 3, 1, D(13000), D("95.38")), ("Pres 2", 3, 1, D(11000), D("99.09"))])
            self.assertEqual([(n["ad"], n["emir"], n["eksik"], n["gecikme"]) for n in r["nedenler"]],
                             [("Mekanik arıza", 2, D(1100), 2), ("Takım kırılması", 1, D(150), 1), ("Kalite reddi", 1, D(100), 0), ("Malzeme gecikmesi", 1, D(0), 2), (main.YOK, 1, D(0), 0)])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]], [("Yüksek", "Gecikmiş açık iş emri", "IE-1008"), ("Orta", "Sapma nedeni yazılmamış", main.YOK)])
            self.assertEqual(load_workbook(cikti).sheetnames, ["İş Emirleri", "Özet", "Sapma Nedenleri", "Hatlar", "Ürünler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "e.csv"
            y.write_text("Üretim Emri No;Mamul;İş Merkezi;Plan Miktar;Üretilen Miktar;Termin;Bitiş Tarihi;Açıklama\nE1;Ü;H;1.000;950,5;05.10.2026;06.10.2026;Arıza\nE2;;;10;;05.10.2026;;\ne1;;;10;;05.10.2026;;\nE3;;;0;;05.10.2026;;\n"
                         "E4;;;10;;x;;\nE5;;;10;;05.10.2026;y;\nE6;;;10;-1;05.10.2026;;\n;;;10;;05.10.2026;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["urun"], x["hat"], x["plan"], x["gercek"], x["p_bitis"], x["g_bitis"], x["neden"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("E1", "Ü", "H", D(1000), D("950.5"), date(2026, 10, 5), date(2026, 10, 6), "Arıza"), ("E2", "(ürün yok)", "(hat yok)", D(10), D(0), date(2026, 10, 5), None, "")],
                              [("İş emri okunamadı", x) for x in ("e1", "E3", "E4", "E5", "E6")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--emirler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "12.10.2026", "--tolerans", "2,5"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--tolerans", "-1"], ["--tolerans", "100"])], [2] * 3)
            self.assertEqual(main.main(["--emirler", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("İş Emri No;Planlanan Miktar\nA;5\n", "İş Emri No;Planlanan Miktar;Planlanan Bitiş\n;5;05.10.2026\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
