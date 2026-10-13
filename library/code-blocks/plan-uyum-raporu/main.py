"""
Plan Uyum Raporu — Workers / Workless kod bloğu
Üretim Planlama › Üretim Planlama Uzmanı

Planlanan ile gerçekleşen üretimi iş emri bazında karşılaştırır; miktar ve tarih sapmalarını nedenleriyle raporlar.
  Miktar durumu        = Eksik (gerçekleşen < plan × (1 − tolerans)) · Fazla (> plan × (1 + tolerans)) · Tam;   --tolerans %0
  Tarih durumu         = Zamanında (gerçekleşen bitiş ≤ planlanan bitiş) · Geç · açık iş emrinde planlanan bitiş geçmişse Gecikmiş (açık)
  İş emri durumu       = Plana uygun (tam ve zamanında) · Eksik · Geç · Eksik ve geç · Fazla üretim · Gecikmiş (açık) · Devam ediyor
  Plan uyumu %         = plana uygun iş emri ÷ değerlendirilen iş emri × 100     (devam edenler değerlendirmeye girmez)
  Miktar uyumu %       = Σ min(gerçekleşen, plan) ÷ Σ plan × 100                 (fazla üretim uyumu artırmaz)
  Zamanında bitirme %  = zamanında biten iş emri ÷ değerlendirilen iş emri × 100
Sapma nedenini dosyada siz yazarsınız; kod neden atamaz.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"no": ("is emri no", "is emri", "uretim emri no", "siparis no"), "urun": ("urun", "urun adi", "stok kodu", "mamul"), "hat": ("hat", "makine", "is merkezi", "uretim hatti"),
         "plan": ("planlanan miktar", "plan miktar", "planlanan"), "gercek": ("gerceklesen miktar", "gerceklesen", "uretilen miktar"), "p_bitis": ("planlanan bitis", "plan bitis", "planlanan bitis tarihi", "termin"),
         "g_bitis": ("gerceklesen bitis", "bitis tarihi", "gerceklesen bitis tarihi"), "neden": ("sapma nedeni", "neden", "aciklama")}
DURUM = ("Gecikmiş (açık)", "Eksik ve geç", "Eksik", "Geç", "Fazla üretim", "Plana uygun", "Devam ediyor")
DURUM_RENK = {"Gecikmiş (açık)": KIRMIZI, "Eksik ve geç": KIRMIZI, "Eksik": SARI, "Geç": SARI, "Fazla üretim": MAVI, "Plana uygun": YESIL, "Devam ediyor": None}
YOK = "(neden yazılmamış)"
SIFIR = Decimal(0)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("no", "plan", "p_bitis")):
        no = metin(r["no"])
        if not no:
            continue
        plan, gercek, pb = para(r["plan"]), para(r["gercek"]) if metin(r["gercek"]) else SIFIR, tarih(r["p_bitis"])
        if plan is None or plan <= 0 or gercek is None or gercek < 0 or pb is None or katla(no) in gorulen or (metin(r["g_bitis"]) and tarih(r["g_bitis"]) is None):
            uy.append(uyari("Yüksek", "İş emri okunamadı", no, f"Satır {r['_satir']}: " + ("aynı iş emri numarası daha önce geçti" if katla(no) in gorulen else "miktarlardan ya da tarihlerden biri okunamadı") + ". Satır alınmadı."))
            continue
        gorulen.add(katla(no))
        l.append({"no": no, "urun": metin(r["urun"]) or "(ürün yok)", "hat": metin(r["hat"]) or "(hat yok)", "plan": plan, "gercek": gercek, "p_bitis": pb, "g_bitis": tarih(r["g_bitis"]), "neden": metin(r["neden"])})
    return l, uy


def _oran(pay, payda) -> Decimal | None:
    return yuvarla(Decimal(pay) * 100 / Decimal(payda)) if payda else None


def uyum(emirler: list[dict], bugun: date, tolerans: Decimal = SIFIR) -> dict:
    if not emirler:
        raise ValueError("Değerlendirilecek iş emri yok.")
    uy = []
    for e in emirler:
        e["fark"], e["oran"] = e["gercek"] - e["plan"], _oran(e["gercek"], e["plan"])
        e["miktar"] = "Eksik" if e["gercek"] < e["plan"] * (100 - tolerans) / 100 else "Fazla" if e["gercek"] > e["plan"] * (100 + tolerans) / 100 else "Tam"
        e["eksik"], e["kabul"] = max(e["plan"] - e["gercek"], SIFIR), min(e["gercek"], e["plan"])
        if e["g_bitis"] is None:
            e["gecikme"] = (bugun - e["p_bitis"]).days if e["p_bitis"] < bugun else 0
            e["durum"] = "Gecikmiş (açık)" if e["gecikme"] else "Devam ediyor"
            if e["gecikme"]:
                uy.append(uyari("Yüksek", "Gecikmiş açık iş emri", e["no"], f"{e['urun']} · {e['hat']}: planlanan bitiş {e['p_bitis']:%d.%m.%Y} ({e['gecikme']} gün geçti); gerçekleşen {e['gercek']:,.0f} / {e['plan']:,.0f}. "
                                f"Neden: {e['neden'] or YOK}."))
        else:
            e["gecikme"] = max((e["g_bitis"] - e["p_bitis"]).days, 0)
            gec = e["gecikme"] > 0
            e["durum"] = ("Eksik ve geç" if gec else "Eksik") if e["miktar"] == "Eksik" else "Geç" if gec else "Fazla üretim" if e["miktar"] == "Fazla" else "Plana uygun"
        e["zamaninda"] = e["g_bitis"] is not None and e["gecikme"] == 0

    def ozet(l: list[dict]) -> dict:
        d = [e for e in l if e["durum"] != "Devam ediyor"]
        plan, kabul = sum((e["plan"] for e in d), SIFIR), sum((e["kabul"] for e in d), SIFIR)
        return {"emir": len(d), "devam": len(l) - len(d), "uygun": sum(1 for e in d if e["durum"] == "Plana uygun"), "uyum": _oran(sum(1 for e in d if e["durum"] == "Plana uygun"), len(d)), "plan": plan,
                "gercek": sum((e["gercek"] for e in d), SIFIR), "miktar_uyum": _oran(kabul, plan), "eksik": sum((e["eksik"] for e in d), SIFIR), "fazla": sum((max(e["fark"], SIFIR) for e in d), SIFIR),
                "gec": sum(1 for e in d if e["gecikme"] > 0), "zamaninda": _oran(sum(1 for e in d if e["zamaninda"]), len(d))}

    def kirilim(alan: str) -> list[dict]:
        grup = {}
        for e in emirler:
            grup.setdefault(katla(e[alan]), []).append(e)
        return sorted(({"ad": g[0][alan], **ozet(g)} for g in grup.values()), key=lambda s: (s["miktar_uyum"] is None, s["miktar_uyum"] or 0, katla(s["ad"])))

    sapmali = [e for e in emirler if e["durum"] not in ("Plana uygun", "Devam ediyor")]
    neden = {}
    for e in sapmali:
        n = neden.setdefault(katla(e["neden"]), {"ad": e["neden"] or YOK, "emir": 0, "eksik": SIFIR, "gecikme": 0, "emirler": []})
        n["emir"], n["eksik"], n["gecikme"] = n["emir"] + 1, n["eksik"] + e["eksik"], n["gecikme"] + e["gecikme"]
        n["emirler"].append(e["no"])
    nedenler = sorted(neden.values(), key=lambda n: (n["ad"] == YOK, -n["emir"], -n["eksik"], katla(n["ad"])))
    for n in nedenler:
        n["pay"] = _oran(n["emir"], len(sapmali))
        if n["ad"] == YOK:
            uy.append(uyari("Orta", "Sapma nedeni yazılmamış", YOK, f"{n['emir']} sapmalı iş emrinde neden yok: {', '.join(n['emirler'])}."))
    return {"emirler": sorted(emirler, key=lambda e: (DURUM.index(e["durum"]), -e["gecikme"], -e["eksik"], katla(e["no"]))), "ozet": ozet(emirler), "hatlar": kirilim("hat"), "urunler": kirilim("urun"), "nedenler": nedenler,
            "durumlar": [{"durum": d, "emir": sum(1 for e in emirler if e["durum"] == d)} for d in DURUM], "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, bugun: date, tolerans: Decimal) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "İş Emirleri"
    _baslik(ws, ["Durum", "İş Emri No", "Ürün", "Hat", "Planlanan", "Gerçekleşen", "Fark", "Gerçekleşme %", "Planlanan Bitiş", "Gerçekleşen Bitiş", "Gecikme (gün)", "Sapma Nedeni"], (18, 14, 24, 16, 13, 13, 12, 12, 12, 12, 10, 34))
    for e in r["emirler"]:
        _satir(ws, [e["durum"], e["no"], e["urun"], e["hat"], e["plan"], e["gercek"], e["fark"], e["oran"], e["p_bitis"], e["g_bitis"], e["gecikme"] or None, e["neden"]], para_s=(8,), tarih_s=(9, 10),
               dolgu={1: DURUM_RENK[e["durum"]], 7: KIRMIZI if e["miktar"] == "Eksik" else None}, sar=(12,))
    ws = wb.create_sheet("Özet")
    o = r["ozet"]
    _baslik(ws, ["Gösterge", "Değer"], (44, 16))
    for etiket, deger, para_mi in (("Değerlendirilen iş emri", o["emir"], False), ("Devam eden iş emri (değerlendirme dışı)", o["devam"], False), ("Plana uygun iş emri", o["uygun"], False), ("Plan uyumu %", o["uyum"], True),
                                   ("Miktar uyumu %", o["miktar_uyum"], True), ("Zamanında bitirme %", o["zamaninda"], True), ("Planlanan miktar", o["plan"], True), ("Gerçekleşen miktar", o["gercek"], True),
                                   ("Eksik kalan miktar", o["eksik"], True), ("Plan üstü üretim", o["fazla"], True), ("Geç biten ya da gecikmiş iş emri", o["gec"], False)):
        _satir(ws, [etiket, deger], para_s=(2,) if para_mi else ())
    ws.append([])
    _baslik(ws, ["İş emri durumu", "Adet"], dondur="A2")
    for d in r["durumlar"]:
        _satir(ws, [d["durum"], d["emir"]], dolgu={1: DURUM_RENK[d["durum"]] if d["emir"] else None})
    ws.append([])
    ws.append([f"Durum tarihi: {bugun:%d.%m.%Y} · miktar toleransı %{tolerans}"])
    ws = wb.create_sheet("Sapma Nedenleri")
    _baslik(ws, ["Sapma Nedeni", "İş Emri", "Pay %", "Eksik Miktar", "Toplam Gecikme (gün)", "İş Emirleri"], (34, 9, 9, 14, 12, 60))
    for n in r["nedenler"]:
        _satir(ws, [n["ad"], n["emir"], n["pay"], n["eksik"] or None, n["gecikme"] or None, ", ".join(n["emirler"])], para_s=(3,), dolgu={1: SARI if n["ad"] == YOK else None}, sar=(6,))
    for baslik, alan, ad in (("Hatlar", "hatlar", "Hat"), ("Ürünler", "urunler", "Ürün")):
        ws = wb.create_sheet(baslik)
        _baslik(ws, [ad, "Değerlendirilen İş Emri", "Plana Uygun", "Plan Uyumu %", "Planlanan", "Gerçekleşen", "Miktar Uyumu %", "Eksik Miktar", "Plan Üstü", "Geç / Gecikmiş", "Zamanında %", "Devam Eden"], (26, 12, 10, 11, 14, 14, 12, 13, 12, 10, 11, 9))
        for s in r[alan]:
            _satir(ws, [s["ad"], s["emir"], s["uygun"], s["uyum"], s["plan"], s["gercek"], s["miktar_uyum"], s["eksik"] or None, s["fazla"] or None, s["gec"] or None, s["zamaninda"], s["devam"] or None], para_s=(4, 7, 11))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, tolerans: Decimal = SIFIR) -> dict:
    emirler, uy = oku(yol)
    r = uyum(emirler, bugun, tolerans)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, bugun, tolerans)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Planlanan ile gerçekleşen üretimi iş emri bazında karşılaştırır; sapmaları nedenleriyle raporlar.")
    p.add_argument("--emirler", type=Path, help="İş Emri No, Ürün, Hat, Planlanan Miktar, Gerçekleşen Miktar, Planlanan Bitiş, Gerçekleşen Bitiş, Sapma Nedeni")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--tolerans", default="0", help="Miktar toleransı, yüzde; plan ± bu oran içindeki üretim 'tam' sayılır (varsayılan 0)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "plan_uyum_raporu.xlsx")
    a = p.parse_args(argv)
    ornek = a.emirler is None
    yol = ORNEK / "is_emirleri.csv" if ornek else a.emirler
    bugun, tol = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today(), para(a.tolerans)
    if bugun is None or tol is None or not 0 <= tol < 100:
        print("[X] --tarih GG.AA.YYYY; --tolerans 0 ile 100 arasında yüzde olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, bugun, tol)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {o['emir']} iş emri değerlendirildi (devam eden {o['devam']})" + (f" · plan uyumu %{o['uyum']} · miktar uyumu %{o['miktar_uyum']} · zamanında %{o['zamaninda']}" if o["emir"] else "") + f" · eksik {o['eksik']:,.0f}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
