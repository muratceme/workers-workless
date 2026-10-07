"""
SPC Kontrol Grafiği ve Cp/Cpk — Workers / Workless kod bloğu
Üretim › Kalite Mühendisi

Alt grup ölçümlerinden X̄-R kontrol grafiği sınırlarını, kontrol dışı noktaları (Nelson kuralları)
ve şartname sınırları verilirse proses yeterlilik indekslerini (Cp, Cpk, Pp, Ppk) hesaplar.
Sonuçları Excel'e yerleşik X̄ ve R grafikleriyle birlikte yazar. İnternete bağlanmaz.

Kullanım:
    python main.py                                              # örnek ölçümlerle dener
    python main.py --girdi olcumler.xlsx --alt-sinir 9.95 --ust-sinir 10.05
    python main.py --girdi olcumler.csv --ozellik "Çap"          # birden çok özellik varsa
"""
from __future__ import annotations

import argparse
import csv
import math
import statistics
import sys
from collections import OrderedDict
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent

# Kontrol grafiği sabitleri (alt grup büyüklüğü n = 2..10) — ASTM / AIAG SPC tabloları
SABITLER = {
    2: (1.880, 0.000, 3.267, 1.128), 3: (1.023, 0.000, 2.574, 1.693), 4: (0.729, 0.000, 2.282, 2.059),
    5: (0.577, 0.000, 2.114, 2.326), 6: (0.483, 0.000, 2.004, 2.534), 7: (0.419, 0.076, 1.924, 2.704),
    8: (0.373, 0.136, 1.864, 2.847), 9: (0.337, 0.184, 1.816, 2.970), 10: (0.308, 0.223, 1.777, 3.078),
}  # (A2, D3, D4, d2)


def sayi(x) -> float | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return float(x)
    try:
        return float(str(x).strip().replace(",", "."))
    except ValueError:
        return None


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ayirici = csv.Sniffer().sniff(metin.splitlines()[0], delimiters=",;\t").delimiter
        satirlar = list(csv.reader(metin.splitlines(), delimiter=ayirici))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def alt_gruplari_cikar(satirlar: list[list], ozellik: str | None = None) -> "OrderedDict[str, list[float]]":
    """İki biçim desteklenir:
    - Geniş: her satır bir alt grup; 'Ölçüm 1', 'Ölçüm 2'... (veya 'X1'...) sütunları
    - Uzun: 'Alt Grup' ve 'Ölçüm' (veya 'Değer') sütunları; aynı alt grubun ölçümleri ayrı satırlarda"""
    b = [kucuk(x) for x in satirlar[0]]
    i_ozellik = next((i for i, x in enumerate(b) if x in {"özellik", "karakteristik", "ölçü"}), None)
    if ozellik and i_ozellik is None:
        raise SystemExit("'Özellik' sütunu yok; --ozellik kullanılamaz.")
    veri = [r for r in satirlar[1:] if not ozellik or str(r[i_ozellik]).strip() == ozellik]
    i_grup = next((i for i, x in enumerate(b) if x in {"alt grup", "altgrup", "numune no", "grup", "saat", "parti"}), None)
    olcum_sutunlari = [i for i, x in enumerate(b) if x.startswith("ölçüm") or x.startswith("olcum") or (x[:1] == "x" and x[1:].isdigit())]
    gruplar: "OrderedDict[str, list[float]]" = OrderedDict()
    if len(olcum_sutunlari) >= 2:
        for n, r in enumerate(veri, 1):
            ad = str(r[i_grup]).strip() if i_grup is not None else str(n)
            gruplar[ad] = [v for v in (sayi(r[i]) for i in olcum_sutunlari if i < len(r)) if v is not None]
    else:
        i_deger = next((i for i, x in enumerate(b) if x in {"ölçüm", "değer", "olcum", "deger", "x"}), None)
        if i_grup is None or i_deger is None:
            raise SystemExit("Ölçüm sütunları bulunamadı. Geniş biçim (Ölçüm 1, Ölçüm 2...) ya da uzun biçim "
                             f"(Alt Grup + Ölçüm) kullanın. Başlıklar: {satirlar[0]}")
        for r in veri:
            v = sayi(r[i_deger])
            if v is not None:
                gruplar.setdefault(str(r[i_grup]).strip(), []).append(v)
    return gruplar


@dataclass
class Sonuc:
    n: int
    alt_grup_adlari: list
    ortalamalar: list
    araliklar: list
    x_ort: float
    r_ort: float
    x_ust: float
    x_alt: float
    r_ust: float
    r_alt: float
    sigma_ic: float
    sigma_toplam: float
    ihlaller: dict = field(default_factory=dict)       # alt grup indeksi -> [kural]
    r_ihlaller: list = field(default_factory=list)
    yeterlilik: dict = field(default_factory=dict)
    uyarilar: list = field(default_factory=list)


def nelson(degerler: list[float], merkez: float, sigma: float) -> dict[int, list[str]]:
    """Nelson kuralları 1, 2, 3, 5, 6 (X̄ grafiği için)."""
    ihlal: dict[int, list[str]] = {}

    def isaretle(i, kural):
        ihlal.setdefault(i, [])
        if kural not in ihlal[i]:
            ihlal[i].append(kural)

    z = [(v - merkez) / sigma if sigma else 0 for v in degerler]
    for i, zi in enumerate(z):
        if abs(zi) > 3:
            isaretle(i, "K1: 3σ dışında")
    for i in range(8, len(z)):
        pencere = z[i - 8:i + 1]
        if all(x > 0 for x in pencere) or all(x < 0 for x in pencere):
            isaretle(i, "K2: 9 nokta merkezin aynı tarafında")
    for i in range(5, len(degerler)):
        p = degerler[i - 5:i + 1]
        if all(p[k] < p[k + 1] for k in range(5)) or all(p[k] > p[k + 1] for k in range(5)):
            isaretle(i, "K3: 6 nokta sürekli artış/azalış")
    for i in range(2, len(z)):
        p = z[i - 2:i + 1]
        if sum(x > 2 for x in p) >= 2 or sum(x < -2 for x in p) >= 2:
            isaretle(i, "K5: 3 noktanın 2'si 2σ dışında (aynı taraf)")
    for i in range(4, len(z)):
        p = z[i - 4:i + 1]
        if sum(x > 1 for x in p) >= 4 or sum(x < -1 for x in p) >= 4:
            isaretle(i, "K6: 5 noktanın 4'ü 1σ dışında (aynı taraf)")
    return ihlal


def hesapla(gruplar: "OrderedDict[str, list[float]]", alt_sinir: float | None = None, ust_sinir: float | None = None) -> Sonuc:
    boyutlar = {len(v) for v in gruplar.values() if v}
    n = max(boyutlar, key=lambda b: sum(len(v) == b for v in gruplar.values())) if boyutlar else 0
    uyarilar = []
    if n not in SABITLER:
        raise SystemExit(f"Alt grup büyüklüğü 2-10 arasında olmalı (bulunan: {n}).")
    tam = OrderedDict((k, v) for k, v in gruplar.items() if len(v) == n)
    atilan = [k for k, v in gruplar.items() if len(v) != n]
    if atilan:
        uyarilar.append(f"Eksik/fazla ölçümlü {len(atilan)} alt grup hesaba katılmadı: {', '.join(atilan[:10])}")
    if len(tam) < 20:
        uyarilar.append(f"Yalnızca {len(tam)} alt grup var; güvenilir kontrol sınırları için en az 20-25 alt grup önerilir")
    a2, d3, d4, d2 = SABITLER[n]
    ort = [statistics.fmean(v) for v in tam.values()]
    ara = [max(v) - min(v) for v in tam.values()]
    x_ort, r_ort = statistics.fmean(ort), statistics.fmean(ara)
    sigma_ic = r_ort / d2
    tum = [x for v in tam.values() for x in v]
    sigma_top = statistics.stdev(tum) if len(tum) > 1 else 0.0
    s = Sonuc(n=n, alt_grup_adlari=list(tam), ortalamalar=ort, araliklar=ara, x_ort=x_ort, r_ort=r_ort,
              x_ust=x_ort + a2 * r_ort, x_alt=x_ort - a2 * r_ort, r_ust=d4 * r_ort, r_alt=d3 * r_ort,
              sigma_ic=sigma_ic, sigma_toplam=sigma_top, uyarilar=uyarilar)
    s.ihlaller = nelson(ort, x_ort, a2 * r_ort / 3)
    s.r_ihlaller = [i for i, r in enumerate(ara) if r > s.r_ust or r < s.r_alt]

    if alt_sinir is not None or ust_sinir is not None:
        y = {}
        if alt_sinir is not None and ust_sinir is not None:
            if ust_sinir <= alt_sinir:
                raise SystemExit("Üst şartname sınırı alt sınırdan büyük olmalı.")
            y["Cp"] = (ust_sinir - alt_sinir) / (6 * sigma_ic) if sigma_ic else math.inf
            y["Pp"] = (ust_sinir - alt_sinir) / (6 * sigma_top) if sigma_top else math.inf
        cpk = [v for v in ((ust_sinir - x_ort) / (3 * sigma_ic) if ust_sinir is not None and sigma_ic else None,
                           (x_ort - alt_sinir) / (3 * sigma_ic) if alt_sinir is not None and sigma_ic else None) if v is not None]
        ppk = [v for v in ((ust_sinir - x_ort) / (3 * sigma_top) if ust_sinir is not None and sigma_top else None,
                           (x_ort - alt_sinir) / (3 * sigma_top) if alt_sinir is not None and sigma_top else None) if v is not None]
        if cpk:
            y["Cpk"] = min(cpk)
        if ppk:
            y["Ppk"] = min(ppk)
        sinir_disi = sum(1 for x in tum if (alt_sinir is not None and x < alt_sinir) or (ust_sinir is not None and x > ust_sinir))
        y["Şartname dışı ölçüm"] = sinir_disi
        y["Şartname dışı oranı (%)"] = 100 * sinir_disi / len(tum) if tum else 0
        s.yeterlilik = y
        if s.ihlaller or s.r_ihlaller:
            uyarilar.append("Proses istatistiksel kontrolde değil; Cp/Cpk yorumlanmadan önce özel nedenler giderilmeli")
    return s


def yorum(cpk: float | None) -> str:
    if cpk is None:
        return ""
    if cpk >= 1.67:
        return "Çok iyi (≥ 1,67)"
    if cpk >= 1.33:
        return "Yeterli (≥ 1,33)"
    if cpk >= 1.0:
        return "Sınırda (1,00 – 1,33): iyileştirme gerekli"
    return "Yetersiz (< 1,00): şartname dışı üretim bekleniyor"


BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")


def rapor_yaz(s: Sonuc, cikti: Path, baslik: str, alt_sinir=None, ust_sinir=None) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kontrol Grafiği"
    ws.append(["Alt Grup", "X̄", "R", "X̄ UCL", "X̄ CL", "X̄ LCL", "R UCL", "R CL", "R LCL", "Kontrol Dışı (Nelson)"])
    for h in ws[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for i, ad in enumerate(s.alt_grup_adlari):
        kurallar = list(s.ihlaller.get(i, []))
        if i in s.r_ihlaller:
            kurallar.append("R grafiği sınır dışı")
        ws.append([ad, round(s.ortalamalar[i], 6), round(s.araliklar[i], 6), round(s.x_ust, 6), round(s.x_ort, 6),
                   round(s.x_alt, 6), round(s.r_ust, 6), round(s.r_ort, 6), round(s.r_alt, 6), "; ".join(kurallar)])
        if kurallar:
            for h in ws[ws.max_row]:
                h.fill = KIRMIZI
    son = ws.max_row
    for j, g in enumerate((12, 12, 12, 12, 12, 12, 12, 12, 12, 60), 1):
        ws.column_dimensions[get_column_letter(j)].width = g
    ws.freeze_panes = "B2"

    for ad, sutunlar, yer in (("X̄ Grafiği", (2, 4, 5, 6), "L2"), ("R Grafiği", (3, 7, 8, 9), "L22")):
        g = LineChart()
        g.title, g.height, g.width = f"{baslik} · {ad}", 9, 22
        for c in sutunlar:
            g.add_data(Reference(ws, min_col=c, min_row=1, max_row=son), titles_from_data=True)
        g.set_categories(Reference(ws, min_col=1, min_row=2, max_row=son))
        for seri in g.series[1:]:
            seri.smooth = False
            seri.graphicalProperties.line.dashStyle = "dash"
        ws.add_chart(g, yer)

    oz = wb.create_sheet("Özet")
    satirlar = [
        ["Özellik", baslik], ["Alt grup büyüklüğü (n)", s.n], ["Alt grup sayısı", len(s.alt_grup_adlari)],
        ["X̄̄ (genel ortalama)", s.x_ort], ["R̄ (ortalama aralık)", s.r_ort],
        ["X̄ UCL / LCL", f"{s.x_ust:.6g} / {s.x_alt:.6g}"], ["R UCL / LCL", f"{s.r_ust:.6g} / {s.r_alt:.6g}"],
        ["σ (alt grup içi, R̄/d2)", s.sigma_ic], ["σ (toplam, örneklem std)", s.sigma_toplam],
        ["Kontrol dışı alt grup", len(set(s.ihlaller) | set(s.r_ihlaller))],
    ]
    if alt_sinir is not None or ust_sinir is not None:
        satirlar.append(["Şartname (LSL / USL)", f"{alt_sinir if alt_sinir is not None else '-'} / {ust_sinir if ust_sinir is not None else '-'}"])
    for k, v in s.yeterlilik.items():
        satirlar.append([k, round(v, 4) if isinstance(v, float) and math.isfinite(v) else v])
    if "Cpk" in s.yeterlilik:
        satirlar.append(["Cpk yorumu", yorum(s.yeterlilik["Cpk"])])
    for u in s.uyarilar:
        satirlar.append(["Uyarı", u])
    for r in satirlar:
        oz.append(r)
    oz.column_dimensions["A"].width = 32
    oz.column_dimensions["B"].width = 90
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description="X̄-R kontrol grafiği, Nelson kuralları ve Cp/Cpk.")
    ap.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "mil_capi.csv")
    ap.add_argument("--ozellik", help="Girdide 'Özellik' sütunu varsa hangisi analiz edilecek")
    ap.add_argument("--alt-sinir", type=float, help="Alt şartname sınırı (LSL)")
    ap.add_argument("--ust-sinir", type=float, help="Üst şartname sınırı (USL)")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "spc.xlsx")
    a = ap.parse_args(argv)
    if a.girdi == BURASI / "ornek_veri" / "mil_capi.csv" and a.alt_sinir is None and a.ust_sinir is None:
        a.alt_sinir, a.ust_sinir = 24.95, 25.05      # örnek: Ø25 ±0,05 mm
    gruplar = alt_gruplari_cikar(oku(a.girdi), a.ozellik)
    s = hesapla(gruplar, a.alt_sinir, a.ust_sinir)
    rapor_yaz(s, a.cikti, a.ozellik or a.girdi.stem, a.alt_sinir, a.ust_sinir)
    print(f"[OK] {len(s.alt_grup_adlari)} alt grup (n={s.n}) · X̄̄={s.x_ort:.5g} · R̄={s.r_ort:.5g}")
    if s.yeterlilik:
        print("[OK] " + " · ".join(f"{k}={v:.3f}" for k, v in s.yeterlilik.items() if k in {"Cp", "Cpk", "Pp", "Ppk"}))
    print(f"[{'!' if s.ihlaller or s.r_ihlaller else 'OK'}] Kontrol dışı alt grup: {len(set(s.ihlaller) | set(s.r_ihlaller))}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
