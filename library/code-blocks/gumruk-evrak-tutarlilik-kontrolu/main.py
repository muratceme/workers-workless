"""
Gümrük Evrak Tutarlılık Kontrolü — Workers / Workless kod bloğu
Lojistik ve Taşımacılık › Gümrük Operasyon › Gümrük Operasyon Uzmanı

Fatura, çeki listesi, menşe belgesi ve taşıma belgesindeki bilgileri karşılaştırıp uyumsuzlukları listeler:
  - Belge bilgileri (gönderici, alıcı, fatura no, kap adedi, brüt/net ağırlık, varış ülkesi, menşe…) belgeler
    arasında aynı mı; sayısal alanlarda tolerans (--tolerans, varsayılan %0,5).
  - Kalem bazında: fatura ↔ çeki listesi (miktar), fatura ↔ menşe belgesi (miktar, GTİP'in ilk 6 hanesi, menşe
    ülke); bir belgede olup diğerinde olmayan kalemler.
  - Fatura içi: miktar × birim fiyat = tutar, kalem toplamı = fatura toplamı.
  - Çeki listesi: net ≤ brüt, kalem toplamları ↔ belge bilgilerindeki kap ve ağırlık.
  - Teslim şekli Incoterms 2020 kurallarından biri mi; yalnız deniz/iç su yolu için olan FAS, FOB, CFR, CIF'in
    karayolu/havayolu/demiryolu belgesiyle kullanılması.
  - GTİP biçimi (Türkiye GTİP 12 hane; 6 / 8 haneli kodlar bilgi olarak işaretlenir).
İnternete bağlanmaz.

Kullanım:
    python main.py                                                     # örnek ihracat evrakı (bilerek hatalı)
    python main.py --klasor ./sevkiyat_evraki --tolerans 0.01
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
SIFIR = Decimal(0)
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
INCOTERMS = {"EXW", "FCA", "CPT", "CIP", "DAP", "DPU", "DDP", "FAS", "FOB", "CFR", "CIF"}
DENIZ_YOLU = {"FAS", "FOB", "CFR", "CIF"}
DENIZ_BELGELERI = ("konsimento", "bill of lading", "b l", "bl", "sea waybill", "deniz")
SAYISAL_ALANLAR = ("kap", "agirlik", "toplam", "tutar", "hacim", "miktar")
TARIH_ALANLARI = ("tarih",)
HUKUKI_EKLER = r"\b(a s|as|ltd|sti|gmbh|inc|llc|co|corp|sa|srl|bv|ag|limited|anonim|sirketi)\b"
BELGE_DOSYALARI = {"Fatura": ("fatura",), "Çeki Listesi": ("ceki", "packing"), "Menşe Belgesi": ("mense", "origin", "eur 1", "eur1", "atr"),
                   "Taşıma Belgesi": ("tasima", "cmr", "konsimento", "awb")}
KALEM_SUTUNLARI = {"kod": ("urun kodu", "kod", "stok kodu", "malzeme kodu"), "tanim": ("tanim", "urun", "aciklama", "mal tanimi"),
                   "gtip": ("gtip", "hs kodu", "hs code", "gtip no"), "miktar": ("miktar", "adet"), "birim": ("birim",),
                   "fiyat": ("birim fiyat", "fiyat"), "tutar": ("tutar", "toplam"), "mense": ("mense", "mense ulke", "mense ulkesi"),
                   "kap": ("kap adedi", "kap", "koli"), "net": ("net agirlik kg", "net agirlik", "net kg"), "brut": ("brut agirlik kg", "brut agirlik", "brut kg")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def unvan_katla(s) -> str:
    return " ".join(re.sub(HUKUKI_EKLER, " ", katla(s)).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(str(x or "").strip(), f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif s.count(".") > 1 or (s.count(".") == 1 and len(s.split(".")[1]) == 3):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def sayi(x) -> str:
    if x is None:
        return "—"
    s = f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return s[:-3] if s.endswith(",00") else s


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def gtip_rakam(x) -> str:
    return re.sub(r"\D", "", str(x or ""))


# ----------------------------------------------------------------------------
# Girdiler
# ----------------------------------------------------------------------------

@dataclass
class Kalem:
    belge: str
    satir: int
    kod: str
    tanim: str
    gtip: str
    miktar: Decimal | None
    birim: str
    fiyat: Decimal | None
    tutar: Decimal | None
    mense: str
    kap: Decimal | None
    net: Decimal | None
    brut: Decimal | None

    @property
    def anahtar(self) -> str:
        return katla(self.kod) or katla(self.tanim)


def kalemleri_oku(yol: Path, belge: str) -> list[Kalem]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in KALEM_SUTUNLARI.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if ("kod" in es or "tanim" in es) and "miktar" in es:
            break
    else:
        raise ValueError(f"{yol.name}: kalem başlığı bulunamadı (Ürün Kodu / Tanım ve Miktar gerekli)")
    al = lambda r, a: r[es[a]] if a in es and es[a] < len(r) else None  # noqa: E731
    sonuc = []
    for n, r in enumerate(s[bi + 1:], bi + 2):
        if not (al(r, "kod") or al(r, "tanim")) or katla(al(r, "tanim")).startswith("toplam"):
            continue
        sonuc.append(Kalem(belge, n, str(al(r, "kod") or "").strip(), str(al(r, "tanim") or "").strip(), str(al(r, "gtip") or "").strip(),
                           para(al(r, "miktar")), str(al(r, "birim") or "").strip(), para(al(r, "fiyat")), para(al(r, "tutar")),
                           str(al(r, "mense") or "").strip(), para(al(r, "kap")), para(al(r, "net")), para(al(r, "brut"))))
    return sonuc


def bilgileri_oku(yol: Path) -> tuple[list[str], dict[str, dict[str, str]]]:
    s = tablo_oku(yol)
    belgeler = [str(c).strip() for c in s[0][1:] if c]
    alanlar = {}
    for r in s[1:]:
        if r and r[0]:
            alanlar[str(r[0]).strip()] = {b: str(r[j + 1]).strip() for j, b in enumerate(belgeler) if j + 1 < len(r) and r[j + 1] not in (None, "")}
    return belgeler, alanlar


def belge_turu(ad: str) -> str | None:
    k = katla(ad)
    return next((b for b, anahtarlar in BELGE_DOSYALARI.items() if any(a in k for a in anahtarlar)), None)


def klasoru_oku(klasor: Path) -> tuple[list[str], dict, dict[str, list[Kalem]], list[str]]:
    belgeler, alanlar, kalemler, atlanan = [], {}, {}, []
    for yol in sorted(klasor.iterdir()):
        if yol.suffix.lower() not in {".csv", ".xlsx", ".xlsm"}:
            atlanan.append(yol.name)
            continue
        if "bilgi" in katla(yol.stem):
            belgeler, alanlar = bilgileri_oku(yol)
            continue
        tur = belge_turu(yol.stem)
        if tur is None:
            atlanan.append(f"{yol.name} (belge türü adından anlaşılamadı)")
            continue
        try:
            kalemler[tur] = kalemleri_oku(yol, tur)
        except ValueError as h:
            atlanan.append(str(h))
    return belgeler, alanlar, kalemler, atlanan


# ----------------------------------------------------------------------------
# Kontroller
# ----------------------------------------------------------------------------

def alan_bul(alanlar: dict, *anahtarlar: str) -> tuple[str, dict] | tuple[None, dict]:
    for ad, d in alanlar.items():
        if any(a in katla(ad) for a in anahtarlar):
            return ad, d
    return None, {}


def kontrol_et(belgeler: list[str], alanlar: dict, kalemler: dict[str, list[Kalem]], tolerans: Decimal) -> tuple[list[dict], dict]:
    bulgular, durum = [], defaultdict(dict)    # durum[alan][belge] = "Uyumsuz" vb.

    def b(onem, konu, aciklama, belge=""):
        bulgular.append({"onem": onem, "konu": konu, "belge": belge, "aciklama": aciklama})

    for ad, d in alanlar.items():
        if len(d) < 2:
            continue
        k = katla(ad)
        if any(x in k for x in SAYISAL_ALANLAR):
            degerler = {bl: para(v) for bl, v in d.items()}
            gecerli = {bl: v for bl, v in degerler.items() if v is not None}
            if not gecerli:
                continue
            sayac = defaultdict(int)
            for v in gecerli.values():
                sayac[v] += 1
            ref = max(sayac, key=lambda v: (sayac[v], v))          # çoğunluğun değeri; eşitlikte büyük olan
            for bl, v in gecerli.items():
                if ref and abs(ref - v) / ref > tolerans:
                    durum[ad][bl] = "Uyumsuz"
            farkli = [bl for bl in gecerli if durum[ad].get(bl)]
            if farkli:
                b("Yüksek" if "kap" in k else "Orta", ad, " · ".join(f"{bl}: {sayi(v)}" for bl, v in gecerli.items()) +
                  f" (fark {sayi(max(gecerli.values()) - min(gecerli.values()))}; tolerans %{sayi(tolerans * 100)})", ", ".join(farkli))
        elif any(x in k for x in TARIH_ALANLARI):
            tarihler = {bl: tarih(v) for bl, v in d.items()}
            if len({t for t in tarihler.values() if t}) > 1:
                for bl in d:
                    durum[ad][bl] = "Uyumsuz"
                b("Orta", ad, " · ".join(f"{bl}: {v}" for bl, v in d.items()), ", ".join(d))
        else:
            norm = unvan_katla if any(x in k for x in ("gonderici", "alici", "ihracatci", "ithalatci")) else katla
            gruplar = defaultdict(list)
            for bl, v in d.items():
                gruplar[norm(v)].append(bl)
            if len(gruplar) > 1:
                cogunluk = max(gruplar.values(), key=len)
                for bl in d:
                    if bl not in cogunluk or len(cogunluk) == 1:
                        durum[ad][bl] = "Uyumsuz"
                b("Yüksek" if any(x in k for x in ("fatura no", "alici", "gonderici")) else "Orta", ad, " · ".join(f"{bl}: {v}" for bl, v in d.items()),
                  ", ".join(bl for bl in d if durum[ad].get(bl)))

    _, teslim = alan_bul(alanlar, "teslim sekli", "incoterm")
    _, tasima_turu = alan_bul(alanlar, "tasima belgesi turu", "belge turu")
    for bl, v in teslim.items():
        kod = re.sub(r"[^A-Z]", " ", v.upper()).split()
        kod = next((x for x in kod if x in INCOTERMS), None)
        if kod is None:
            b("Orta", "Teslim şekli", f"'{v}' Incoterms 2020 kurallarından biri değil ({', '.join(sorted(INCOTERMS))})", bl)
        elif kod in DENIZ_YOLU:
            tur = katla(" ".join(tasima_turu.values()))
            if tur and not any(x in f" {tur} " for x in (f" {d} " for d in DENIZ_BELGELERI)):
                b("Orta", "Teslim şekli", f"{kod} yalnız deniz ve iç su yolu taşımacılığı içindir; taşıma belgesi türü '{' / '.join(tasima_turu.values())}'. "
                                          "Diğer taşıma şekillerinde FCA, CPT, CIP gibi her taşıma şekline uygun kurallar kullanılır", bl)

    fat = kalemler.get("Fatura", [])
    for k_ in fat:
        if k_.miktar is not None and k_.fiyat is not None and k_.tutar is not None:
            hesap = (k_.miktar * k_.fiyat).quantize(Decimal("0.01"))
            if abs(hesap - k_.tutar) > Decimal("0.01"):
                b("Yüksek", "Fatura hesap", f"{k_.kod or k_.tanim}: {sayi(k_.miktar)} × {sayi(k_.fiyat)} = {sayi(hesap)}, faturada {sayi(k_.tutar)}", "Fatura")
        g = gtip_rakam(k_.gtip)
        if g and len(g) != 12:
            b("Bilgi", "GTİP", f"{k_.kod or k_.tanim}: GTİP {k_.gtip} {len(g)} hane (Türkiye GTİP 12 hanedir; beyannamede tamamlayın)", "Fatura")
    ad_t, toplam = alan_bul(alanlar, "fatura toplami", "toplam tutar")
    if fat and toplam.get("Fatura"):
        kt = sum((k_.tutar or SIFIR for k_ in fat), SIFIR)
        if abs(kt - (para(toplam["Fatura"]) or SIFIR)) > Decimal("0.01"):
            durum[ad_t]["Fatura"] = "Uyumsuz"
            b("Yüksek", "Fatura toplamı", f"Kalem tutarları toplamı {sayi(kt)}, fatura toplamı {toplam['Fatura']}", "Fatura")

    ceki = kalemler.get("Çeki Listesi", [])
    for k_ in ceki:
        if k_.net is not None and k_.brut is not None and k_.net > k_.brut:
            b("Yüksek", "Çeki ağırlık", f"{k_.kod or k_.tanim}: net {sayi(k_.net)} kg brüt {sayi(k_.brut)} kg'dan büyük", "Çeki Listesi")
    for alan_adi, attr, anahtar in (("Kap adedi", "kap", "kap"), ("Net ağırlık", "net", "net agirlik"), ("Brüt ağırlık", "brut", "brut agirlik")):
        if not ceki or all(getattr(k_, attr) is None for k_ in ceki):
            continue
        top = sum((getattr(k_, attr) or SIFIR for k_ in ceki), SIFIR)
        ad, d = alan_bul(alanlar, anahtar)
        if d.get("Çeki Listesi") and para(d["Çeki Listesi"]) is not None and abs(top - para(d["Çeki Listesi"])) > max(top * tolerans, Decimal("0.01")):
            durum[ad]["Çeki Listesi"] = "Uyumsuz"
            b("Orta", alan_adi, f"Çeki listesi kalem toplamı {sayi(top)}, belge bilgisinde {d['Çeki Listesi']}", "Çeki Listesi")

    kalem_satirlari = []
    if fat:
        diger = {bl: {k_.anahtar: k_ for k_ in kl} for bl, kl in kalemler.items() if bl != "Fatura"}
        fmap = {k_.anahtar: k_ for k_ in fat}
        for k_ in fat:
            satir = {"kod": k_.kod, "tanim": k_.tanim, "Fatura": k_.miktar, "notlar": []}
            for bl, m in diger.items():
                x = m.get(k_.anahtar)
                if x is None:
                    satir[bl] = None
                    satir["notlar"].append(f"{bl}nde yok")
                    b("Yüksek" if bl != "Taşıma Belgesi" else "Orta", "Eksik kalem", f"{k_.kod or k_.tanim} faturada var, {bl.lower()}nde yok", bl)
                    continue
                satir[bl] = x.miktar
                if x.miktar is not None and k_.miktar is not None and x.miktar != k_.miktar:
                    satir["notlar"].append(f"{bl} miktarı farklı")
                    b("Yüksek", "Miktar", f"{k_.kod or k_.tanim}: faturada {sayi(k_.miktar)}, {bl.lower()}nde {sayi(x.miktar)}", bl)
                if x.gtip and k_.gtip and gtip_rakam(x.gtip)[:6] != gtip_rakam(k_.gtip)[:6]:
                    satir["notlar"].append(f"{bl} GTİP farklı")
                    b("Yüksek", "GTİP", f"{k_.kod or k_.tanim}: faturada {k_.gtip}, {bl.lower()}nde {x.gtip} (ilk 6 hane farklı)", bl)
                if x.mense and k_.mense and katla(x.mense) != katla(k_.mense):
                    satir["notlar"].append(f"{bl} menşe farklı")
                    b("Yüksek", "Menşe", f"{k_.kod or k_.tanim}: faturada {k_.mense}, {bl.lower()}nde {x.mense}", bl)
            kalem_satirlari.append(satir)
        for bl, m in diger.items():
            for an, x in m.items():
                if an not in fmap:
                    b("Yüksek", "Fazla kalem", f"{x.kod or x.tanim} {bl.lower()}nde var, faturada yok", bl)
                    kalem_satirlari.append({"kod": x.kod, "tanim": x.tanim, "Fatura": None, bl: x.miktar, "notlar": ["Faturada yok"]})
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    bulgular.sort(key=lambda x: sira[x["onem"]])
    return bulgular, {"durum": durum, "kalem_satirlari": kalem_satirlari}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, YESIL, KONTROL = PatternFill("solid", fgColor="FDE2E1"), PatternFill("solid", fgColor="E3F4E1"), PatternFill("solid", fgColor="FFF4CE")
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def rapor_yaz(cikti: Path, belgeler, alanlar, kalemler, bulgular, ek, atlanan) -> None:
    wb = Workbook()
    o = wb.active
    o.title = "Bulgular"
    _baslik(o, ["Önem", "Konu", "Belge", "Açıklama", "Düzeltme / Not"], (9, 18, 16, 90, 30))
    for x in bulgular:
        o.append([x["onem"], x["konu"], x["belge"], x["aciklama"], ""])
        o.cell(o.max_row, 1).fill = PatternFill("solid", fgColor=RENK[x["onem"]])
        o.cell(o.max_row, 5).fill = KONTROL
        o.cell(o.max_row, 4).alignment = UST
    if not bulgular:
        o.append(["—", "", "", "Uyumsuzluk bulunmadı", ""])
    for a in atlanan:
        o.append(["Bilgi", "Okunamayan", a, "", ""])

    bb = wb.create_sheet("Belge Bilgileri")
    _baslik(bb, ["Alan", *belgeler], (24, *[24] * len(belgeler)))
    for ad, d in alanlar.items():
        bb.append([ad, *[d.get(bl, "") for bl in belgeler]])
        for j, bl in enumerate(belgeler, 2):
            if ek["durum"].get(ad, {}).get(bl):
                bb.cell(bb.max_row, j).fill = KIRMIZI
            elif bl in d and len(d) > 1:
                bb.cell(bb.max_row, j).fill = YESIL

    kk = wb.create_sheet("Kalem Karşılaştırma")
    diger = [bl for bl in BELGE_DOSYALARI if bl in kalemler and bl != "Fatura"]
    _baslik(kk, ["Ürün Kodu", "Tanım", "Fatura Miktar", *[f"{bl} Miktar" for bl in diger], "Not"], (12, 34, 13, *[16] * len(diger), 40))
    for s in ek["kalem_satirlari"]:
        kk.append([s["kod"], s["tanim"], s.get("Fatura") and float(s["Fatura"]), *[s.get(bl) and float(s[bl]) for bl in diger], "; ".join(s["notlar"])])
        if s["notlar"]:
            kk.cell(kk.max_row, 4 + len(diger)).fill = KIRMIZI

    for bl in (x for x in BELGE_DOSYALARI if x in kalemler):
        kl = kalemler[bl]
        ws = wb.create_sheet(bl[:31])
        _baslik(ws, ["Satır", "Ürün Kodu", "Tanım", "GTİP", "Miktar", "Birim", "Birim Fiyat", "Tutar", "Menşe", "Kap", "Net kg", "Brüt kg"],
                (6, 12, 34, 18, 10, 7, 10, 12, 10, 6, 9, 9))
        for k_ in kl:
            ws.append([k_.satir, k_.kod, k_.tanim, k_.gtip, *(v and float(v) for v in (k_.miktar,)), k_.birim,
                       *(v and float(v) for v in (k_.fiyat, k_.tutar)), k_.mense, *(v and float(v) for v in (k_.kap, k_.net, k_.brut))])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(klasor: Path, cikti: Path, tolerans: Decimal = Decimal("0.005")) -> dict:
    belgeler, alanlar, kalemler, atlanan = klasoru_oku(klasor)
    if not alanlar and not kalemler:
        raise ValueError(f"{klasor}: belge bilgileri veya kalem dosyası bulunamadı")
    bulgular, ek = kontrol_et(belgeler, alanlar, kalemler, tolerans)
    rapor_yaz(cikti, belgeler, alanlar, kalemler, bulgular, ek, atlanan)
    return {"bulgular": bulgular, "belgeler": belgeler, "alanlar": alanlar, "kalemler": kalemler, "atlanan": atlanan, **ek}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Fatura, çeki listesi, menşe ve taşıma belgesindeki bilgileri karşılaştırıp uyumsuzlukları listeler.")
    p.add_argument("--klasor", type=Path, default=ORNEK, help="Sevkiyat evrakı klasörü (belge_bilgileri + kalem dosyaları)")
    p.add_argument("--tolerans", default="0.005", help="Sayısal alanlarda göreli tolerans (varsayılan 0.005 = %%0,5)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "gumruk_evrak_kontrolu.xlsx")
    a = p.parse_args(argv)
    if not a.klasor.is_dir():
        print(f"[X] Klasör bulunamadı: {a.klasor}")
        return 1
    try:
        s = calistir(a.klasor, a.cikti, Decimal(a.tolerans))
    except (ValueError, InvalidOperation) as h:
        print(f"[X] {h}")
        return 1
    say = defaultdict(int)
    for x in s["bulgular"]:
        say[x["onem"]] += 1
    print(f"[OK] {len(s['kalemler'])} kalem belgesi, {len(s['alanlar'])} belge alanı · {len(s['bulgular'])} bulgu "
          f"(Yüksek {say['Yüksek']}, Orta {say['Orta']}, Bilgi {say['Bilgi']})")
    for x in s["bulgular"]:
        if x["onem"] == "Yüksek":
            print(f"[X] {x['konu']}: {x['aciklama']}")
    for x in s["atlanan"]:
        print(f"[!] {x}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
