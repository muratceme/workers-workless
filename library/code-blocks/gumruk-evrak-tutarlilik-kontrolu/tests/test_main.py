import shutil
import sys
import tempfile
import unittest
from collections import Counter
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "g.xlsx"
        cls.s = main.calistir(ORNEK, cls.cikti)
        cls.b = {(x["konu"], x["belge"]): x["aciklama"] for x in cls.s["bulgular"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_sayilar(self):
        self.assertEqual(Counter(x["onem"] for x in self.s["bulgular"]), {"Yüksek": 7, "Orta": 2})
        self.assertEqual(set(self.s["kalemler"]), {"Fatura", "Çeki Listesi", "Menşe Belgesi"})

    def test_belge_alanlari(self):
        self.assertIn(("Alıcı", "Menşe Belgesi"), self.b)                         # "GmbH." farkı sayılmaz
        self.assertIn(("Fatura No", "Taşıma Belgesi"), self.b)
        self.assertIn(("Kap Adedi", "Fatura"), self.b)
        self.assertIn("fark 60", self.b[("Brüt Ağırlık (kg)", "Taşıma Belgesi")])
        self.assertNotIn(("Fatura Tarihi", "Fatura, Çeki Listesi"), self.b)

    def test_kalemler(self):
        self.assertEqual(self.b[("Fatura hesap", "Fatura")], "MT-4060: 1.200 × 1,90 = 2.280, faturada 2.380")
        self.assertEqual(self.b[("Miktar", "Çeki Listesi")], "MT-4060: faturada 1.200, çeki listesinde 1.150")
        self.assertIn("HT-70140: faturada 6302.60.00.00.00, menşe belgesinde 6302.31", self.b[("GTİP", "Menşe Belgesi")])
        self.assertEqual(self.b[("Eksik kalem", "Menşe Belgesi")], "PS-5050 faturada var, menşe belgesinde yok")

    def test_incoterms(self):
        self.assertIn("FOB yalnız deniz ve iç su yolu", self.b[("Teslim şekli", "Fatura")])

    def test_tolerans(self):
        s = main.calistir(ORNEK, self.cikti, D("0.05"))
        self.assertFalse(any(x["konu"].startswith("Brüt") for x in s["bulgular"]))
        self.assertTrue(any(x["konu"] == "Kap Adedi" for x in s["bulgular"]) is False)   # 1/96 < %5

    def test_excel(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Bulgular", "Belge Bilgileri", "Kalem Karşılaştırma", "Fatura", "Çeki Listesi", "Menşe Belgesi"])
        bb = wb["Belge Bilgileri"]
        satir = next(r for r in bb.iter_rows(min_row=2) if r[0].value == "Alıcı")
        self.assertEqual(satir[3].fill.fgColor.rgb[-6:], "FDE2E1")
        self.assertEqual(satir[1].fill.fgColor.rgb[-6:], "E3F4E1")


class KuralTesti(unittest.TestCase):
    def test_tutarli_evrak_ve_incoterms(self):
        with tempfile.TemporaryDirectory() as tmp:
            k = Path(tmp) / "e"
            k.mkdir()
            (k / "belge_bilgileri.csv").write_text("Alan;Fatura;Taşıma Belgesi\nTeslim Şekli;FCA İstanbul;\nTaşıma Belgesi Türü;;Konşimento\n"
                                                   "Alıcı;ABC Ltd.;ABC Limited\n", encoding="utf-8")
            (k / "fatura.csv").write_text("Ürün Kodu;Miktar;Birim Fiyat;Tutar;GTİP\nA;10;2,5;25;630260\n", encoding="utf-8")
            (k / "ceki.csv").write_text("Ürün Kodu;Miktar;Net Ağırlık;Brüt Ağırlık\nA;10;12;11\n", encoding="utf-8")
            (k / "notlar.txt").write_text("x", encoding="utf-8")
            s = main.calistir(k, Path(tmp) / "o.xlsx")
            konular = [(x["konu"], x["onem"]) for x in s["bulgular"]]
            self.assertEqual(konular, [("Çeki ağırlık", "Yüksek"), ("GTİP", "Bilgi")])
            self.assertIn("notlar.txt", s["atlanan"])
            (k / "belge_bilgileri.csv").write_text("Alan;Fatura\nTeslim Şekli;FOT\n", encoding="utf-8")
            self.assertIn("Incoterms 2020 kurallarından biri değil", main.calistir(k, Path(tmp) / "o.xlsx")["bulgular"][1]["aciklama"])

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--klasor", str(Path(tmp) / "yok")]), 1)


if __name__ == "__main__":
    unittest.main()
