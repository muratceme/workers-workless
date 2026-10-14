"""
Zincir Market Listeleme Takibi — Workers / Workless kod bloğu
Satış ve Kanal Yönetimi › Kanal Satış Yöneticisi
Zincir marketlerde listelenen ürünleri, raf fiyatlarını ve stoksuz mağazaları takip eder.
  Güncel durum     = her zincir · mağaza · ürün için en son ziyaret kaydı
  Listeleme oranı  = ürünün listeli olduğu mağaza ÷ ziyaret edilen mağaza × 100   (zincir bazında)
  Stoksuzluk (OOS) = listeli olup rafta stok olmayan mağaza ÷ listeli mağaza × 100
  Fiyat sapması    = (ortalama raf fiyatı − önerilen raf fiyatı) ÷ önerilen × 100
  Uyarı            = listeleme --liste (%80) altında · stoksuzluk --oos (%10) üstünde · fiyat sapması ± --fiyat (%5) dışında
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

TARAMA = {"tarih": ("tarih", "ziyaret tarihi"), "zincir": ("zincir", "zincir adi", "musteri", "zincir market"), "magaza": ("magaza", "magaza adi", "sube", "magaza kodu"),
          "kod": ("urun kodu", "stok kodu", "barkod", "kod"), "listeli": ("listeli", "listelenmis", "listeleme", "rafta listeli"), "fiyat": ("raf fiyati", "fiyat", "etiket fiyati"),
          "stok": ("stok durumu", "stok", "rafta stok", "stok adedi", "raf stogu")}
URUN = {"kod": ("urun kodu", "stok kodu", "barkod", "kod"), "ad": ("urun adi", "urun", "aciklama"), "tsf": ("onerilen raf fiyati", "tavsiye satis fiyati", "tsf", "onerilen fiyat")}
EVET, HAYIR = {"evet", "e", "var", "x", "1", "listeli", "mevcut"}, {"hayir", "h", "yok", "0", "listesiz", "degil", "listeli degil"}
SIFIR = Decimal(0)


def _evet(x) -> bool | None:
    k = katla(x)
    return True if k in EVET else False if k in HAYIR else None


def oku_tarama(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, TARAMA, ("zincir", "magaza", "kod")):
        z, m, kod = metin(r["zincir"]), metin(r["magaza"]), metin(r["kod"])
        if not (z or m or kod):
            continue
        t = tarih(r["tarih"]) if metin(r["tarih"]) else None
        fiyat = para(r["fiyat"]) if metin(r["fiyat"]) else None
        listeli = _evet(r["listeli"]) if metin(r["listeli"]) else (True if fiyat or metin(r["stok"]) else None)
        sk = metin(r["stok"])
        stok = None if not sk else (para(sk) > 0 if para(sk) is not None else _evet(sk))
        kim = f"{z} · {m} · {kod}"
        if not (z and m and kod) or (metin(r["tarih"]) and t is None) or listeli is None or (metin(r["fiyat"]) and (fiyat is None or fiyat <= 0)) or (sk and stok is None):
            uy.append(uyari("Yüksek", "Tarama satırı okunamadı", kim, f"Satır {r['_satir']}: zincir, mağaza, ürün, tarih, listeli, fiyat ya da stok bilgisi okunamadı. Satır alınmadı."))
            continue
        l.append({"tarih": t, "zincir": z, "magaza": m, "kod": kod, "listeli": listeli, "fiyat": fiyat, "stok": stok, "satir": r["_satir"]})
    return l, uy


def oku_urun(yol: Path) -> tuple[dict, list[dict]]:
    uy, d = [], {}
    for r in kayitlar(yol, URUN, ("kod",)):
        kod = metin(r["kod"])
        if not kod:
            continue
        tsf = para(r["tsf"]) if metin(r["tsf"]) else None
        if metin(r["tsf"]) and (tsf is None or tsf <= 0):
            uy.append(uyari("Yüksek", "Ürün satırı okunamadı", kod, f"Satır {r['_satir']}: önerilen raf fiyatı okunamadı."))
            tsf = None
        d[katla(kod)] = {"kod": kod, "ad": metin(r["ad"]), "tsf": tsf}
    return d, uy


def takip(tarama: list[dict], urunler: dict | None = None, liste: Decimal = Decimal(80), oos: Decimal = Decimal(10), fiyat: Decimal = Decimal(5)) -> dict:
    if not tarama:
        raise ValueError("Tarama kaydı yok.")
    urunler = urunler or {}
    uy, son = [], {}
    for t in sorted(tarama, key=lambda t: (t["tarih"] or date.min, t["satir"])):
        son[(katla(t["zincir"]), katla(t["magaza"]), katla(t["kod"]))] = t
    gruplar = defaultdict(list)
    for t in son.values():
        gruplar[(katla(t["zincir"]), katla(t["kod"]))].append(t)
    ozet, detay = [], []
    for (zk, kk), l in sorted(gruplar.items()):
        u = urunler.get(kk, {})
        listeli = [t for t in l if t["listeli"]]
        stoksuz = [t for t in listeli if t["stok"] is False]
        fiyatlar = [t["fiyat"] for t in listeli if t["fiyat"]]
        ort = sum(fiyatlar, SIFIR) / len(fiyatlar) if fiyatlar else None
        tsf = u.get("tsf")
        o = {"zincir": l[0]["zincir"], "kod": l[0]["kod"], "ad": u.get("ad", ""), "ziyaret": len(l), "listeli": len(listeli), "liste_oran": yuvarla(Decimal(len(listeli)) * 100 / len(l)),
             "stoksuz": len(stoksuz), "oos": yuvarla(Decimal(len(stoksuz)) * 100 / len(listeli)) if listeli else None, "ort": yuvarla(ort) if ort is not None else None,
             "min": min(fiyatlar) if fiyatlar else None, "max": max(fiyatlar) if fiyatlar else None, "tsf": tsf, "sapma": yuvarla((ort - tsf) * 100 / tsf) if ort is not None and tsf else None,
             "son": max((t["tarih"] for t in l if t["tarih"]), default=None)}
        ozet.append(o)
        kim = f"{o['zincir']} · {o['kod']}"
        if o["liste_oran"] < liste:
            eksik = ", ".join(t["magaza"] for t in sorted(l, key=lambda t: katla(t["magaza"])) if not t["listeli"])
            uy.append(uyari("Orta", "Listeleme eksik", kim, f"{o['listeli']}/{o['ziyaret']} mağazada listeli (%{o['liste_oran']}; hedef %{liste}). Listesiz: {eksik}."))
        if o["oos"] is not None and o["oos"] > oos:
            uy.append(uyari("Orta", "Stoksuz mağaza oranı yüksek", kim, f"Listeli {o['listeli']} mağazadan {o['stoksuz']} tanesinde rafta stok yok (%{o['oos']}; eşik %{oos}): "
                                                                     + ", ".join(t["magaza"] for t in stoksuz) + "."))
        if o["sapma"] is not None and abs(o["sapma"]) > fiyat:
            uy.append(uyari("Orta", "Raf fiyatı önerilenden farklı", kim, f"Ortalama raf fiyatı {o['ort']:,.2f}, önerilen {tsf:,.2f} (sapma %{o['sapma']}; bant ±%{fiyat})."))
        for t in l:
            sap = yuvarla((t["fiyat"] - tsf) * 100 / tsf) if t["fiyat"] and tsf else None
            detay.append({**t, "ad": o["ad"], "tsf": tsf, "sapma": sap})
    for kk in sorted(set(urunler) - {k for _, k in gruplar}):
        uy.append(uyari("Bilgi", "Taramada olmayan ürün", urunler[kk]["kod"], "Ürün listesinde var ama hiçbir mağazada taranmamış."))
    detay.sort(key=lambda t: (katla(t["zincir"]), katla(t["magaza"]), katla(t["kod"])))
    zincirler = []
    for z in sorted({o["zincir"] for o in ozet}, key=katla):
        l = [o for o in ozet if o["zincir"] == z]
        top_z, top_l, top_s = sum(o["ziyaret"] for o in l), sum(o["listeli"] for o in l), sum(o["stoksuz"] for o in l)
        zincirler.append({"zincir": z, "urun": len(l), "magaza": len({katla(t["magaza"]) for t in detay if t["zincir"] == z}), "liste_oran": yuvarla(Decimal(top_l) * 100 / top_z),
                          "oos": yuvarla(Decimal(top_s) * 100 / top_l) if top_l else None})
    return {"ozet": ozet, "detay": detay, "zincirler": zincirler, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, liste: Decimal, oos: Decimal, fiyat: Decimal) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Zincir × Ürün"
    _baslik(ws, ["Zincir", "Ürün Kodu", "Ürün Adı", "Ziyaret Edilen Mağaza", "Listeli", "Listeleme %", "Stoksuz", "Stoksuzluk %", "Ort. Raf Fiyatı", "En Düşük", "En Yüksek", "Önerilen",
                 "Fiyat Sapması %", "Son Ziyaret"], (18, 12, 24, 10, 8, 10, 8, 10, 11, 9, 9, 9, 10, 11), "D2")
    for o in r["ozet"]:
        _satir(ws, [o["zincir"], o["kod"], o["ad"], o["ziyaret"], o["listeli"], o["liste_oran"], o["stoksuz"], o["oos"], o["ort"], o["min"], o["max"], o["tsf"], o["sapma"], o["son"]],
               para_s=(6, 8, 9, 10, 11, 12, 13), tarih_s=(14,),
               dolgu={6: KIRMIZI if o["liste_oran"] < liste else YESIL, 8: None if o["oos"] is None else KIRMIZI if o["oos"] > oos else YESIL,
                      13: None if o["sapma"] is None else SARI if abs(o["sapma"]) > fiyat else None})
    ws = wb.create_sheet("Zincirler")
    _baslik(ws, ["Zincir", "Ürün", "Mağaza", "Listeleme %", "Stoksuzluk %"], (20, 7, 8, 11, 11))
    for z in r["zincirler"]:
        _satir(ws, [z["zincir"], z["urun"], z["magaza"], z["liste_oran"], z["oos"]], para_s=(4, 5))
    ws = wb.create_sheet("Mağaza Detayı")
    _baslik(ws, ["Zincir", "Mağaza", "Ürün Kodu", "Ürün Adı", "Ziyaret", "Listeli", "Stok", "Raf Fiyatı", "Önerilen", "Sapma %"], (18, 18, 12, 24, 11, 8, 8, 10, 9, 9), "C2")
    for t in r["detay"]:
        _satir(ws, [t["zincir"], t["magaza"], t["kod"], t["ad"], t["tarih"], "Evet" if t["listeli"] else "Hayır", None if t["stok"] is None else "Var" if t["stok"] else "Yok", t["fiyat"], t["tsf"], t["sapma"]],
               para_s=(8, 9, 10), tarih_s=(5,), dolgu={6: None if t["listeli"] else KIRMIZI, 7: KIRMIZI if t["listeli"] and t["stok"] is False else None,
                                                       10: SARI if t["sapma"] is not None and abs(t["sapma"]) > fiyat else None})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(tarama_yol: Path, cikti: Path, urun_yol: Path | None = None, liste: Decimal = Decimal(80), oos: Decimal = Decimal(10), fiyat: Decimal = Decimal(5)) -> dict:
    tarama, uy = oku_tarama(tarama_yol)
    urunler, uy2 = oku_urun(urun_yol) if urun_yol else ({}, [])
    r = takip(tarama, urunler, liste, oos, fiyat)
    r["uyarilar"] = sirala(uy + uy2 + r["uyarilar"])
    rapor_yaz(cikti, r, liste, oos, fiyat)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Zincir marketlerde listelenen ürünleri, raf fiyatlarını ve stoksuz mağazaları takip eder.")
    p.add_argument("--tarama", type=Path, help="Tarih, Zincir, Mağaza, Ürün Kodu, Listeli (Evet / Hayır), Raf Fiyatı, Stok Durumu (Var / Yok ya da adet)")
    p.add_argument("--urunler", type=Path, help="İsteğe bağlı: Ürün Kodu, Ürün Adı, Önerilen Raf Fiyatı")
    p.add_argument("--liste", default="80", help="Listeleme oranı bu yüzdenin altındaysa uyarılır (varsayılan 80)")
    p.add_argument("--oos", default="10", help="Stoksuz mağaza oranı bu yüzdenin üstündeyse uyarılır (varsayılan 10)")
    p.add_argument("--fiyat", default="5", help="Ortalama raf fiyatı önerilen fiyattan bu yüzdeden fazla saparsa uyarılır (varsayılan 5)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "zincir_listeleme_takibi.xlsx")
    a = p.parse_args(argv)
    ornek = a.tarama is None
    ty = a.tarama or ORNEK / "magaza_ziyaretleri.csv"
    uy = a.urunler or (ORNEK / "urun_listesi.csv" if ornek else None)
    sayilar = [para(x) for x in (a.liste, a.oos, a.fiyat)]
    if any(x is None or not 0 <= x <= 100 for x in sayilar):
        print("[X] --liste, --oos ve --fiyat 0-100 arası olmalı")
        return 2
    for y in (ty, uy):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(ty, a.cikti, uy, *sayilar)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(r['zincirler'])} zincir · {len(r['ozet'])} zincir × ürün · {len(r['detay'])} mağaza × ürün")
    for z in r["zincirler"]:
        print(f"     {z['zincir']}: {z['magaza']} mağaza · listeleme %{z['liste_oran']} · stoksuzluk %{z['oos']}")
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
