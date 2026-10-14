import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def t(gun, z, m, kod, listeli, fiyat=None, stok=True, satir=0):
    return {"tarih": date(2026, 10, gun), "zincir": z, "magaza": m, "kod": kod, "listeli": listeli, "fiyat": None if fiyat is None else D(fiyat), "stok": stok if listeli else None, "satir": satir}


URUN = {"u1": {"kod": "U1", "ad": "Ürün 1", "tsf": D(100)}, "u2": {"kod": "U2", "ad": "", "tsf": D(50)}, "u3": {"kod": "U3", "ad": "", "tsf": None}}


def veri():
    return [t(10, "Z1", "M1", "U1", True, 100), t(10, "Z1", "M2", "U1", True, 105, False), t(10, "Z1", "M3", "U1", False), t(10, "Z1", "M1", "U2", True, 49),
            t(10, "Z2", "M9", "U1", True, 90), t(1, "Z2", "M9", "u1", True, 100)]


class TakipTesti(unittest.TestCase):
    def test_ozet(self):
        r = main.takip(veri(), URUN)
        self.assertEqual([(o["zincir"], o["kod"], o["ziyaret"], o["listeli"], o["liste_oran"], o["stoksuz"], o["oos"], o["ort"], o["sapma"]) for o in r["ozet"]],
                         [("Z1", "U1", 3, 2, D("66.67"), 1, D("50.00"), D("102.50"), D("2.50")), ("Z1", "U2", 1, 1, D("100.00"), 0, D("0.00"), D("49.00"), D("-2.00")),
                          ("Z2", "U1", 1, 1, D("100.00"), 0, D("0.00"), D("90.00"), D("-10.00"))])
        self.assertEqual([(z["zincir"], z["urun"], z["magaza"], z["liste_oran"], z["oos"]) for z in r["zincirler"]], [("Z1", 2, 3, D("75.00"), D("33.33")), ("Z2", 1, 1, D("100.00"), D("0.00"))])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Bilgi", "Taramada olmayan ürün", "U3"), ("Orta", "Listeleme eksik", "Z1 · U1"), ("Orta", "Raf fiyatı önerilenden farklı", "Z2 · U1"),
                          ("Orta", "Stoksuz mağaza oranı yüksek", "Z1 · U1")])
        self.assertEqual(len(r["detay"]), 5)

    def test_esikler(self):
        r = main.takip(veri(), URUN, D(50), D(60), D(15))
        self.assertEqual([u["tur"] for u in r["uyarilar"]], ["Taramada olmayan ürün"])
        self.assertEqual([main._evet(x) for x in ("Evet", "yok", "Listeli değil", "?")], [True, False, False, None])
        with self.assertRaises(ValueError):
            main.takip([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "magaza_ziyaretleri.csv", cikti, ORNEK / "urun_listesi.csv")
            self.assertEqual([(z["zincir"], z["liste_oran"], z["oos"]) for z in r["zincirler"]], [("Zincir A", D("100.00"), D("5.00")), ("Zincir B", D("63.33"), D("21.05")), ("Zincir C", D("76.00"), D("10.53"))])
            tur = [u["tur"] for u in r["uyarilar"]]
            self.assertEqual((tur.count("Listeleme eksik"), tur.count("Raf fiyatı önerilenden farklı"), tur.count("Stoksuz mağaza oranı yüksek"), len(tur)), (3, 6, 6, 15))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Zincir × Ürün", "Zincirler", "Mağaza Detayı", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, z = Path(tmp) / "t.csv", Path(tmp) / "u.csv"
            y.write_text("Zincir;Şube;Barkod;Listeli;Fiyat;Raf Stoğu\nA;1;X;Evet;12,50;0\nA;2;X;;10;\nA;3;X;belki;;\nA;4;X;Evet;-1;\nA;5;X;Hayır;;\n;6;X;Evet;;\n", encoding="utf-8")
            z.write_text("Barkod;TSF\nX;12\nY;x\n", encoding="utf-8")
            l, uy = main.oku_tarama(y)
            self.assertEqual(([(x["magaza"], x["listeli"], x["fiyat"], x["stok"]) for x in l], len(uy)), ([("1", True, D("12.50"), False), ("2", True, D(10), None), ("5", False, None, None)], 3))
            d, uy = main.oku_urun(z)
            self.assertEqual((d["x"]["tsf"], d["y"]["tsf"], len(uy)), (D(12), None, 1))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--tarama", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--urunler", str(z), "--liste", "90", "--oos", "5", "--fiyat", "3"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--liste", "x"], ["--oos", "101"])], [2] * 2)
            self.assertEqual(main.main(arg + ["--urunler", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Zincir;Fiyat\nA;5\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
