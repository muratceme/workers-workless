import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def a(bas, bit=None, tur="Çek", on_ek="A", basamak=4):
    return {"tur": tur, "on_ek": on_ek, "bas": bas, "bit": bit if bit is not None else bas, "basamak": basamak, "sayan": "", "satir": 2}


def f(r):
    return [(x["tur"], x["fark"], x["bas"], x["bit"], x["adet"]) for x in r["farklar"]]


class SeriTesti(unittest.TestCase):
    def test_coz_ve_yaz(self):
        self.assertEqual([main.seri_coz(x) for x in ("A 0012501", "a0012501", "TM-000301", "12345", 12345, 12345.0, "ç 07", "A", "", None, "12A", "A1B2")],
                         [("A", 12501, 7), ("A", 12501, 7), ("TM", 301, 6), ("", 12345, 5), ("", 12345, 5), ("", 12345, 5), ("C", 7, 2), None, None, None, None, ("A1B", 2, 1)])
        self.assertEqual([main.seri_yaz(*x) for x in (("A", 12501, 7), ("", 45, 5), ("TM", 1234567, 6))], ["A 0012501", "00045", "TM 1234567"])
        self.assertEqual(main._araliklar({("A", 3), ("A", 1), ("A", 2), ("A", 5), ("B", 6), ("B", 7), ("", 9)}), [("", 9, 9), ("A", 1, 3), ("A", 5, 5), ("B", 6, 7)])


class KarsilastirTesti(unittest.TestCase):
    def test_farklar(self):
        r = main.karsilastir([a(1, 10)], [a(1, 3), a(5, 6), a(9, 12), a(11, 11), a(20)])
        self.assertEqual(f(r), [("Çek", "Eksik", "A 0004", "A 0004", 1), ("Çek", "Eksik", "A 0007", "A 0008", 2), ("Çek", "Fazla", "A 0011", "A 0012", 2), ("Çek", "Fazla", "A 0020", "A 0020", 1), ("Çek", "Mükerrer sayım", "A 0011", "A 0011", 1)])
        self.assertEqual((r["ozet"], r["toplam"], r["say"]), ([{"tur": "Çek", "kayitli": 10, "sayilan": 10, "eksik": 3, "fazla": 3, "mukerrer": 1, "sonuc": "Fark var"}],
                                                              {"kayitli": 10, "sayilan": 10, "eksik": 3, "fazla": 3, "mukerrer": 1}, {"Fark var": 1, "Kontrol et": 0, "Tam": 0}))
        self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                         [("Yüksek", "Eksik kıymetli evrak", "Çek · A 0004"), ("Yüksek", "Eksik kıymetli evrak", "Çek · A 0007 – A 0008"), ("Yüksek", "Kayıt dışı kıymetli evrak", "Çek · A 0011 – A 0012"),
                          ("Yüksek", "Kayıt dışı kıymetli evrak", "Çek · A 0020"), ("Orta", "Mükerrer sayım", "Çek · A 0011")])

    def test_tur_ve_on_ek(self):
        kayit = [a(1, 5), a(1, 5, on_ek="B"), a(4, 8), a(1, 3, "Kart", "", 6), a(1, 2, "Form", "F")]
        sayim = [a(1, 8), a(1, 5, on_ek="b".upper()), a(1, 3, "KART", "", 6), a(1, 2, "Senet", "S")]
        r = main.karsilastir(kayit, sayim)
        self.assertEqual([(o["tur"], o["kayitli"], o["sayilan"], o["eksik"], o["fazla"], o["mukerrer"], o["sonuc"]) for o in r["ozet"]],
                         [("Çek", 13, 13, 0, 0, 0, "Kontrol et"), ("Form", 2, 0, 2, 0, 0, "Fark var"), ("Kart", 3, 3, 0, 0, 0, "Tam"), ("Senet", 0, 2, 0, 2, 0, "Fark var")])
        self.assertEqual(f(r), [("Çek", "Mükerrer kayıt", "A 0004", "A 0005", 2), ("Form", "Eksik", "F 0001", "F 0002", 2), ("Senet", "Fazla", "S 0001", "S 0002", 2)])
        self.assertEqual(r["say"], {"Fark var": 2, "Kontrol et": 1, "Tam": 1})
        r = main.karsilastir([a(1, 5)], [a(1, 5, on_ek="B")])
        self.assertEqual(f(r), [("Çek", "Eksik", "A 0001", "A 0005", 5), ("Çek", "Fazla", "B 0001", "B 0005", 5)])
        self.assertEqual((main.karsilastir([a(1, 3)], [a(1, 2), a(2, 3)])["ozet"][0]["sonuc"], main.karsilastir([a(1, 3)], [])["toplam"]["eksik"], main.karsilastir([], [a(7)])["toplam"]["fazla"]), ("Kontrol et", 3, 1))
        with self.assertRaises(ValueError):
            main.karsilastir([], [])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "kayit.csv", ORNEK / "sayim.csv", cikti)
            self.assertEqual([(o["tur"], o["kayitli"], o["sayilan"], o["eksik"], o["fazla"], o["mukerrer"], o["sonuc"]) for o in r["ozet"]],
                             [("Banka Kartı (kişiselleştirilmemiş)", 150, 155, 0, 5, 0, "Fark var"), ("Bloke Çek", 20, 19, 1, 0, 0, "Fark var"), ("Çek Karnesi (10 yaprak)", 60, 60, 0, 0, 0, "Tam"),
                              ("Çek Karnesi (25 yaprak)", 40, 38, 2, 0, 0, "Fark var"), ("Teminat Mektubu Formu", 50, 50, 0, 0, 5, "Kontrol et")])
            self.assertEqual((r["toplam"], r["say"]), ({"kayitli": 320, "sayilan": 322, "eksik": 3, "fazla": 5, "mukerrer": 5}, {"Fark var": 3, "Kontrol et": 1, "Tam": 1}))
            self.assertEqual(f(r), [("Banka Kartı (kişiselleştirilmemiş)", "Fazla", "K 100151", "K 100155", 5), ("Bloke Çek", "Eksik", "BC 0000120", "BC 0000120", 1), ("Çek Karnesi (25 yaprak)", "Eksik", "B 0004026", "B 0004027", 2),
                                    ("Teminat Mektubu Formu", "Mükerrer sayım", "TM 000336", "TM 000340", 5)])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Yüksek", "Eksik kıymetli evrak", "Bloke Çek · BC 0000120"), ("Yüksek", "Eksik kıymetli evrak", "Çek Karnesi (25 yaprak) · B 0004026 – B 0004027"),
                              ("Yüksek", "Kayıt dışı kıymetli evrak", "Banka Kartı (kişiselleştirilmemiş) · K 100151 – K 100155"), ("Orta", "Mükerrer sayım", "Teminat Mektubu Formu · TM 000336 – TM 000340")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Farklar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            ky, sy = Path(tmp) / "k.csv", Path(tmp) / "s.csv"
            ky.write_text("Evrak;Başlangıç;Bitiş\nÇek;a 0001;0010\nÇek;A0020;\nÇek;A 0030;A 0025\nÇek;A 0001;B 0005\nÇek;A 1;A 200000\nÇek;A;A 5\n;A 1;\nÇek;;\nÇek;A 1;x\n", encoding="utf-8")
            sy.write_text("Evrak Türü;Seri No;Son Seri No;Sayan\nÇek;A 0001;A 0010;Ali\n", encoding="utf-8")
            l, uy = main.oku(ky)
            self.assertEqual(([(x["tur"], x["on_ek"], x["bas"], x["bit"], x["basamak"]) for x in l], [u["tur"] for u in uy]), ([("Çek", "A", 1, 10, 4), ("Çek", "A", 20, 20, 4)], ["Satır okunamadı"] * 7))
            self.assertEqual([x["sayan"] for x in main.oku(sy)[0]], ["Ali"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--kayit", str(ky), "--sayim", str(sy), "--cikti", c]
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(["--kayit", str(ky)]), 2)
            self.assertEqual(main.main(["--kayit", str(ky), "--sayim", str(Path(tmp) / "yok.csv")]), 1)
            sy.write_text("Evrak Türü;Sayan\nÇek;Ali\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            ky.write_text("Evrak Türü;Seri Başlangıç\n;\n", encoding="utf-8")
            sy.write_text("Evrak Türü;Seri Başlangıç\n;\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
