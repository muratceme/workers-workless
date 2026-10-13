"""
Kıymetli Evrak Sayım Kontrolü — Workers / Workless kod bloğu
Şube Bankacılığı › Şube Operasyon Yöneticisi

Kıymetli evrak (çek karnesi, kart, teminat mektubu formu) sayımını kayıtlarla karşılaştırır; seri numarası farklarını listeler.
  Seri numarası   = baştaki harf / işaret (ön ek) + sondaki rakamlar; "A 0012501" → ön ek A, numara 12501
  Aralık          = Seri Başlangıç – Seri Bitiş arasındaki bütün numaralar (bitiş boşsa tek numara)
  Eksik           = kayıtta olup sayımda olmayan seri numaraları
  Fazla           = sayımda olup kayıtta olmayan seri numaraları (kayıt dışı)
  Mükerrer sayım  = sayım listesinde birden çok kez geçen seri numaraları
  Sonuç           = Fark var (eksik ya da fazla) · Kontrol et (yalnız mükerrer) · Tam
Farklar ardışık numaralar birleştirilerek aralık olarak yazılır. Karşılaştırma evrak türü ve ön ek bazındadır.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
from collections import Counter

SUTUN = {"tur": ("evrak turu", "evrak", "kiymet turu", "tur"), "bas": ("seri baslangic", "baslangic seri no", "baslangic", "seri no", "ilk seri no"), "bit": ("seri bitis", "bitis seri no", "bitis", "son seri no"),
         "sayan": ("sayan", "sayim yapan", "sorumlu"), "nota": ("not", "aciklama")}
AZAMI = 100000
SONUC_RENK = {"Fark var": KIRMIZI, "Kontrol et": SARI, "Tam": YESIL}
FARK_RENK = {"Eksik": KIRMIZI, "Fazla": SARI, "Mükerrer sayım": MAVI, "Mükerrer kayıt": MAVI}


def seri_coz(s) -> tuple[str, int, int] | None:
    """'A 0012501' → ('A', 12501, 7): ön ek, numara, basamak sayısı."""
    if isinstance(s, float) and s == int(s):
        s = int(s)
    m = re.fullmatch(r"(.*?)(\d+)", re.sub(r"\s+", "", str(s if s is not None else "")))
    return (m.group(1).translate(_TR).upper().strip("-./"), int(m.group(2)), len(m.group(2))) if m else None


def seri_yaz(on_ek: str, n: int, basamak: int) -> str:
    return f"{on_ek} {n:0{basamak}d}".strip()


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("tur", "bas")):
        tur = metin(r["tur"])
        if not tur and not metin(r["bas"]):
            continue
        bas, bit = seri_coz(r["bas"]), seri_coz(r["bit"]) if metin(r["bit"]) else None
        hata = None
        if not tur or bas is None or (metin(r["bit"]) and bit is None):
            hata = "evrak türü ve seri numarası okunamadı (seri numarası rakamla bitmeli)"
        elif bit and bit[0] not in ("", bas[0]):
            hata = f"başlangıç ve bitişin ön eki farklı ({bas[0] or 'yok'} / {bit[0]})"
        elif bit and bit[1] < bas[1]:
            hata = "bitiş numarası başlangıçtan küçük"
        elif bit and bit[1] - bas[1] + 1 > AZAMI:
            hata = f"aralık {AZAMI:,} numaradan büyük; yazım hatası olabilir"
        if hata:
            uy.append(uyari("Yüksek", "Satır okunamadı", tur or f"Satır {r['_satir']}", f"{yol.name} satır {r['_satir']}: '{metin(r['bas'])}' – '{metin(r['bit'])}': {hata}. Satır alınmadı."))
            continue
        l.append({"tur": tur, "on_ek": bas[0], "bas": bas[1], "bit": bit[1] if bit else bas[1], "basamak": bas[2], "sayan": metin(r["sayan"]), "satir": r["_satir"]})
    return l, uy


def _araliklar(anahtarlar) -> list[tuple[str, int, int]]:
    sonuc = []
    for on_ek, n in sorted(anahtarlar):
        if sonuc and sonuc[-1][0] == on_ek and sonuc[-1][2] == n - 1:
            sonuc[-1] = (on_ek, sonuc[-1][1], n)
        else:
            sonuc.append((on_ek, n, n))
    return sonuc


def karsilastir(kayit: list[dict], sayim: list[dict]) -> dict:
    if not kayit and not sayim:
        raise ValueError("Kayıt ve sayım listesi boş.")
    turler = {}
    for kaynak, liste in (("kayit", kayit), ("sayim", sayim)):
        for x in liste:
            t = turler.setdefault(katla(x["tur"]), {"ad": x["tur"], "kayit": Counter(), "sayim": Counter(), "basamak": {}})
            t["basamak"].setdefault(x["on_ek"], x["basamak"])
            t[kaynak].update((x["on_ek"], n) for n in range(x["bas"], x["bit"] + 1))
    uy, ozet, farklar = [], [], []
    for _, t in sorted(turler.items()):
        k, s = t["kayit"], t["sayim"]
        gruplar = (("Eksik", k.keys() - s.keys()), ("Fazla", s.keys() - k.keys()), ("Mükerrer sayım", {a for a, n in s.items() if n > 1}), ("Mükerrer kayıt", {a for a, n in k.items() if n > 1}))
        for ad, kume in gruplar:
            for on_ek, n1, n2 in _araliklar(kume):
                b = t["basamak"][on_ek]
                f = {"tur": t["ad"], "fark": ad, "bas": seri_yaz(on_ek, n1, b), "bit": seri_yaz(on_ek, n2, b), "adet": n2 - n1 + 1}
                farklar.append(f)
                kim = f"{t['ad']} · {f['bas']}" + (f" – {f['bit']}" if n2 > n1 else "")
                if ad == "Eksik":
                    uy.append(uyari("Yüksek", "Eksik kıymetli evrak", kim, f"{f['adet']} adet kayıtta var, sayımda yok. Yeniden sayın; bulunamazsa kurum içi prosedürünüzdeki tutanak ve bildirim adımlarını uygulayın."))
                elif ad == "Fazla":
                    uy.append(uyari("Yüksek", "Kayıt dışı kıymetli evrak", kim, f"{f['adet']} adet sayımda var, kayıtta yok. Giriş kaydı yapılmamış ya da çıkışı yapılmış evrak kasada kalmış olabilir."))
                elif ad == "Mükerrer sayım":
                    uy.append(uyari("Orta", "Mükerrer sayım", kim, f"{f['adet']} seri numarası sayım listesinde birden çok kez geçiyor. Aralıklar çakışıyor; sayım tutanağını düzeltin."))
                else:
                    uy.append(uyari("Orta", "Mükerrer kayıt", kim, f"{f['adet']} seri numarası kayıt listesinde birden çok kez geçiyor."))
        eksik, fazla, mukerrer = len(gruplar[0][1]), len(gruplar[1][1]), len(gruplar[2][1])
        ozet.append({"tur": t["ad"], "kayitli": len(k), "sayilan": len(s), "eksik": eksik, "fazla": fazla, "mukerrer": mukerrer, "sonuc": "Fark var" if eksik or fazla else "Kontrol et" if mukerrer or gruplar[3][1] else "Tam"})
    toplam = {a: sum(o[a] for o in ozet) for a in ("kayitli", "sayilan", "eksik", "fazla", "mukerrer")}
    return {"ozet": ozet, "toplam": toplam, "farklar": farklar, "say": {s: sum(1 for o in ozet if o["sonuc"] == s) for s in SONUC_RENK}, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    _baslik(ws, ["Evrak Türü", "Kayıtlı Adet", "Sayılan Adet", "Eksik", "Fazla", "Mükerrer Sayım", "Sonuç"], (40, 12, 12, 9, 9, 11, 14))
    for o in r["ozet"]:
        _satir(ws, [o["tur"], o["kayitli"], o["sayilan"], o["eksik"] or None, o["fazla"] or None, o["mukerrer"] or None, o["sonuc"]], dolgu={7: SONUC_RENK[o["sonuc"]]})
    t = r["toplam"]
    n = _satir(ws, ["TOPLAM", t["kayitli"], t["sayilan"], t["eksik"], t["fazla"], t["mukerrer"], None])
    for h in ws[n]:
        h.font = Font(bold=True)
    ws = wb.create_sheet("Farklar")
    _baslik(ws, ["Evrak Türü", "Fark", "Seri Başlangıç", "Seri Bitiş", "Adet", "Açıklama (siz doldurun)"], (40, 16, 18, 18, 8, 40))
    for f in r["farklar"]:
        _satir(ws, [f["tur"], f["fark"], f["bas"], f["bit"], f["adet"], None], dolgu={2: FARK_RENK[f["fark"]]})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(k_yol: Path, s_yol: Path, cikti: Path) -> dict:
    kayit, uy1 = oku(k_yol)
    sayim, uy2 = oku(s_yol)
    r = karsilastir(kayit, sayim)
    r["uyarilar"] = sirala(uy1 + uy2 + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Kıymetli evrak sayımını kayıtlarla karşılaştırıp seri numarası farklarını listeler.")
    p.add_argument("--kayit", type=Path, help="Kayıtlara göre mevcut: Evrak Türü, Seri Başlangıç, Seri Bitiş")
    p.add_argument("--sayim", type=Path, help="Sayım listesi: Evrak Türü, Seri Başlangıç, Seri Bitiş, Sayan")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kiymetli_evrak_sayim_kontrolu.xlsx")
    a = p.parse_args(argv)
    if (a.kayit is None) != (a.sayim is None):
        print("[X] --kayit ve --sayim birlikte verilmelidir")
        return 2
    k_yol, s_yol = (ORNEK / "kayit.csv", ORNEK / "sayim.csv") if a.kayit is None else (a.kayit, a.sayim)
    for y in (k_yol, s_yol):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(k_yol, s_yol, a.cikti)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {len(r['ozet'])} evrak türü · kayıtlı {t['kayitli']}, sayılan {t['sayilan']} · eksik {t['eksik']}, fazla {t['fazla']}, mükerrer sayım {t['mukerrer']} · tam {r['say']['Tam']}, fark var {r['say']['Fark var']}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
