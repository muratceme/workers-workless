import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def a(no, dan, genel, tavsiye=None, tur="Bakım", ay=9, geri=False, q=None):
    return {"no": no, "is_emri": "", "tarih": date(2026, ay, 1), "danisman": dan, "tur": tur, "genel": D(genel), "tavsiye": None if tavsiye is None else D(tavsiye),
            "puanlar": {"Karşılama": D(q)} if q is not None else {}, "yorum": "", "geri": geri}


class AnalizTesti(unittest.TestCase):
    def test_oranlar(self):
        l = [a("1", "A", 10, 10, q=9), a("2", "A", 9, 9, q=7), a("3", "A", 6, 5, geri=True), a("4", "A", 8, 8, "Onarım", 8), a("5", "B", 4, 2, "Onarım", 8), a("6", "B", 10)]
        r = main.analiz(l, ["Karşılama"], 10, en_az=3)
        g = r["genel"]
        # NPS: tavsiye 10, 9 → 2 öneren; 5, 2 → 2 kötüleyen; 5 yanıt → 0
        self.assertEqual((g["n"], g["ort"], g["ust"], g["dusuk"], g["nps"], g["nps_n"], g["Karşılama"]), (6, D("7.83"), D("50.00"), 2, D("0.0"), 5, D("8.00")))
        s = {x["ad"]: x for x in r["danisman"]}
        self.assertEqual((s["A"]["ort"], s["A"]["ust"], s["A"]["nps"], s["B"]["nps"], s["B"]["Karşılama"]), (D("8.25"), D("50.00"), D("25.0"), D("-100.0"), None))
        self.assertEqual([(x["ad"], x["n"]) for x in r["ay"]], [("2026-08", 2), ("2026-09", 4)])
        self.assertEqual([x["no"] for x in r["dusukler"]], ["3", "5"])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Bilgi", "Anket sayısı az", "B"), ("Orta", "Düşük puan · geri arama", "5"), ("Yüksek", "Düşük puan · geri arama", "3")])
        r5 = main.analiz([a("1", "A", 5), a("2", "A", 3), a("3", "A", 4)], [], 5)
        self.assertEqual((r5["genel"]["ust"], r5["genel"]["dusuk"], r5["esik"]), (D("33.33"), 1, D(3)))
        self.assertEqual(main.analiz(l, [], 10, D(8))["genel"]["dusuk"], 3)
        with self.assertRaises(ValueError):
            main.analiz([], [])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "anketler.csv", cikti)
            g = r["genel"]
            self.assertEqual((g["n"], g["ort"], g["ust"], g["nps"], g["dusuk"]), (261, D("8.50"), D("52.11"), D("38.9"), 21))
            self.assertEqual(r["sorular"], ["Randevu kolaylığı", "Karşılama", "Açıklama ve fiyat bilgisi", "Zamanında teslim", "İş kalitesi"])
            self.assertEqual(sum(1 for x in r["dusukler"] if x["geri"]), 12)
            self.assertEqual([u["kim"] for u in r["uyarilar"] if u["tur"] in ("Satır okunamadı", "Anket sayısı az")], ["A-7500", "Danışman E (yeni)"])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Danışman", "İş Türü", "Ay", "Geri Arama", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "a.csv"
            y.write_text("Anket No;Danışman;Genel Memnuniyet;Tavsiye;Soru 1: Karşılama;Soru 2 - Teslim;Geri Arama\n1;X;4;3;5;;Evet\n1;X;5;;;;\n2;;11;;;;\n3;;5;11;;;\n4;;5;;x;;\n5;;3,5;;4;5;\n",
                         encoding="utf-8")
            l, s, uy = main.oku(y, 5)
            self.assertEqual(s, ["Karşılama", "Teslim"])
            self.assertEqual(([(x["no"], x["genel"], x["geri"], x["puanlar"]) for x in l], len(uy)),
                             ([("1", D(4), True, {"Karşılama": D(5)}), ("5", D("3.5"), False, {"Karşılama": D(4), "Teslim": D(5)})], 4))
            c = str(Path(tmp) / "x.xlsx")
            self.assertEqual(main.main(["--anket", str(y), "--olcek", "5", "--dusuk", "2", "--min", "1", "--cikti", c]), 0)
            self.assertEqual([main.main(["--anket", str(y), "--cikti", c] + v) for v in (["--olcek", "5", "--dusuk", "5"], ["--dusuk", "x"], ["--min", "0"])], [2] * 3)
            self.assertEqual(main.main(["--anket", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Anket;Puan\n1;5\n", encoding="utf-8")
            self.assertEqual(main.main(["--anket", str(y), "--cikti", c]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
