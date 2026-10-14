"""
Müşteri Memnuniyeti (CSI) Analizi — Workers / Workless kod bloğu
Satış Sonrası Hizmetler (Servis) › Satış Sonrası Hizmetler Müdürü
Servis memnuniyet anketlerini danışman, iş türü ve ay bazında analiz eder; düşük puanlı anketleri geri arama listesine alır.
  Ortalama       = genel memnuniyet puanlarının ortalaması (--olcek 10 ya da 5)
  En üst kutu %  = 10'luk ölçekte 9–10, 5'lik ölçekte 5 veren anket ÷ anket × 100
  Düşük puan     = genel memnuniyet ≤ --dusuk (varsayılan 10'lukta 6, 5'likte 3)
  NPS            = tavsiye 9–10 veren % − tavsiye 0–6 veren % (0–10 tavsiye sorusu varsa)
  Soru puanları  = başlığı "Soru" ile başlayan sütunların ortalaması
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
ALAN = {"no": ("anket no", "anket", "no"), "is_emri": ("is emri no", "is emri"), "tarih": ("tarih", "anket tarihi", "servis tarihi"), "danisman": ("danisman", "servis danismani", "temsilci"),
        "tur": ("is turu", "hizmet turu", "tur"), "genel": ("genel memnuniyet", "genel puan", "memnuniyet", "csi"), "tavsiye": ("tavsiye", "tavsiye puani", "nps"),
        "yorum": ("yorum", "aciklama", "musteri yorumu"), "geri": ("geri arama istiyor", "geri arama", "aranmak istiyor")}
SIFIR = Decimal(0)


def oku(yol: Path, olcek: int = 10) -> tuple[list[dict], list[str], list[dict]]:
    s = tablo_oku(yol)
    bi = next((i for i, r in enumerate(s[:15]) if "genel memnuniyet" in [katla(c) for c in r] or "csi" in [katla(c) for c in r]), None)
    if bi is None:
        raise ValueError(f"{yol.name}: başlık satırı bulunamadı; 'Genel Memnuniyet' sütunu gerekli.")
    b = [katla(c) for c in s[bi]]
    es = {}
    for alan, adlar in ALAN.items():
        j = next((b.index(a) for a in adlar if a in b), None)
        if j is not None:
            es[alan] = j
    sorular = [(j, re.sub(r"^\s*soru\s*\d*\s*[:.\-–]?\s*", "", metin(s[bi][j]), flags=re.I) or metin(s[bi][j])) for j, k in enumerate(b) if k.startswith("soru")]
    uy, l, gorulen = [], [], set()
    for n, r in enumerate(s[bi + 1:], bi + 2):
        al = lambda a: r[es[a]] if a in es and es[a] < len(r) else None
        no = metin(al("no"))
        g = para(al("genel")) if metin(al("genel")) else None
        tv = para(al("tavsiye")) if metin(al("tavsiye")) else None
        t = tarih(al("tarih")) if metin(al("tarih")) else None
        kim = no or f"Satır {n}"
        puanlar = {}
        hata = g is None or not 0 <= g <= olcek or (metin(al("tavsiye")) and (tv is None or not 0 <= tv <= 10)) or (metin(al("tarih")) and t is None)
        for j, ad in sorular:
            v = r[j] if j < len(r) else None
            if metin(v):
                p = para(v)
                if p is None or not 0 <= p <= olcek:
                    hata = True
                puanlar[ad] = p
        if hata:
            uy.append(uyari("Yüksek", "Satır okunamadı", kim, f"Satır {n}: genel memnuniyet (0–{olcek}), tavsiye (0–10), tarih ya da soru puanlarından biri okunamadı / ölçek dışı."))
            continue
        if no and katla(no) in gorulen:
            uy.append(uyari("Orta", "Anket iki kez yazılmış", kim, f"Satır {n}: ikinci satır alınmadı."))
            continue
        if no:
            gorulen.add(katla(no))
        l.append({"no": no, "is_emri": metin(al("is_emri")), "tarih": t, "danisman": metin(al("danisman")) or "Belirtilmemiş", "tur": metin(al("tur")) or "Belirtilmemiş",
                  "genel": g, "tavsiye": tv, "puanlar": puanlar, "yorum": metin(al("yorum")), "geri": katla(al("geri")) in ("evet", "e", "x", "1", "istiyor")})
    return l, [ad for _, ad in sorular], uy


def _ozet(l: list[dict], olcek: int, ust: Decimal, sorular: list[str]) -> dict:
    n = len(l)
    tv = [x["tavsiye"] for x in l if x["tavsiye"] is not None]
    d = {"n": n, "ort": yuvarla(sum((x["genel"] for x in l), SIFIR) / n) if n else None,
         "ust": yuvarla(Decimal(sum(1 for x in l if x["genel"] >= ust)) * 100 / n) if n else None,
         "nps": yuvarla(Decimal(sum(1 for v in tv if v >= 9) - sum(1 for v in tv if v <= 6)) * 100 / len(tv), "0.1") if tv else None, "nps_n": len(tv)}
    for s in sorular:
        p = [x["puanlar"][s] for x in l if x["puanlar"].get(s) is not None]
        d[s] = yuvarla(sum(p, SIFIR) / len(p)) if p else None
    return d


def analiz(anketler: list[dict], sorular: list[str], olcek: int = 10, dusuk: Decimal | None = None, en_az: int = 5) -> dict:
    if not anketler:
        raise ValueError("Okunabilen anket yok.")
    dusuk = dusuk if dusuk is not None else Decimal(6 if olcek == 10 else 3)
    ust = Decimal(9 if olcek == 10 else olcek)
    uy = []
    for a in anketler:
        a["dusuk"] = a["genel"] <= dusuk
        if a["dusuk"]:
            kim = a["no"] or a["is_emri"] or "Anket"
            uy.append(uyari("Yüksek" if a["geri"] else "Orta", "Düşük puan · geri arama", kim,
                            f"{a['danisman']} · {a['tur']}: genel {a['genel']}/{olcek}" + (" · müşteri geri arama istiyor" if a["geri"] else "") + (f" · \"{a['yorum']}\"" if a["yorum"] else "")))

    def kirilim(fn, sirali=False):
        g = {}
        for a in anketler:
            g.setdefault(fn(a), []).append(a)
        satir = [{"ad": k, **_ozet(v, olcek, ust, sorular), "dusuk": sum(1 for a in v if a["dusuk"])} for k, v in g.items() if k is not None]
        return sorted(satir, key=lambda s: s["ad"]) if sirali else sorted(satir, key=lambda s: (-s["n"], katla(s["ad"])))
    danisman = kirilim(lambda a: a["danisman"])
    for s in danisman:
        if s["n"] < en_az:
            uy.append(uyari("Bilgi", "Anket sayısı az", s["ad"], f"{s['n']} anket (en az {en_az}); oranlar karşılaştırma için güvenilir değil."))
    genel = {"ad": "TOPLAM", **_ozet(anketler, olcek, ust, sorular), "dusuk": sum(1 for a in anketler if a["dusuk"])}
    return {"genel": genel, "danisman": danisman, "tur": kirilim(lambda a: a["tur"]), "ay": kirilim(lambda a: f"{a['tarih']:%Y-%m}" if a["tarih"] else None, True),
            "dusukler": sorted((a for a in anketler if a["dusuk"]), key=lambda a: (not a["geri"], a["genel"], a["tarih"] or date.min)), "sorular": sorular, "olcek": olcek,
            "esik": dusuk, "en_az": en_az, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Danışman"
    bas = ["Anket", "Ortalama", "En Üst Kutu %", "Düşük Puan", "NPS", "NPS Yanıtı"] + r["sorular"]
    for ad, l in (("Danışman", r["danisman"]), ("İş Türü", r["tur"]), ("Ay", r["ay"])):
        if ad != "Danışman":
            ws = wb.create_sheet(ad)
        _baslik(ws, [ad] + bas, (22, 8, 10, 9, 8, 9) + (14,) * len(r["sorular"]))
        for s in l + [r["genel"]]:
            n = _satir(ws, [s["ad"], s["n"], s["ort"], s["ust"], s["dusuk"], s["nps"], s["nps_n"]] + [s[q] for q in r["sorular"]], para_s=(3, 4) + tuple(range(8, 8 + len(r["sorular"]))),
                       dolgu={1: MAVI if s["n"] < r["en_az"] and s["ad"] != "TOPLAM" else None})
            if s["ad"] == "TOPLAM":
                for h in ws[n]:
                    h.font = Font(bold=True)
    ws = wb.create_sheet("Geri Arama")
    _baslik(ws, ["Anket No", "İş Emri", "Tarih", "Danışman", "İş Türü", "Genel", "Tavsiye", "Geri Arama İstiyor", "Yorum", "Sonuç / Not"], (10, 12, 11, 18, 14, 7, 7, 9, 50, 30))
    for a in r["dusukler"]:
        _satir(ws, [a["no"], a["is_emri"], a["tarih"], a["danisman"], a["tur"], a["genel"], a["tavsiye"], "Evet" if a["geri"] else "", a["yorum"], ""], tarih_s=(3,), sar=(9,),
               dolgu={8: KIRMIZI if a["geri"] else None})
    ws = wb.active
    ws.append([])
    ws.append([f"Ölçek {r['olcek']} · düşük puan ≤ {r['esik']} · en üst kutu: {'9–10' if r['olcek'] == 10 else r['olcek']} · NPS = %9–10 − %0–6 (tavsiye sorusu)."])
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, olcek: int = 10, dusuk: Decimal | None = None, en_az: int = 5) -> dict:
    anketler, sorular, uy = oku(yol, olcek)
    r = analiz(anketler, sorular, olcek, dusuk, en_az)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Servis memnuniyet anketlerini danışman ve iş türü bazında analiz eder.")
    p.add_argument("--anket", type=Path, help="Anket No, İş Emri No, Tarih, Danışman, İş Türü, Genel Memnuniyet, Tavsiye (0-10), Soru: … sütunları, Yorum, Geri Arama İstiyor")
    p.add_argument("--olcek", type=int, choices=(5, 10), default=10, help="Genel memnuniyet ve soru ölçeği (5 ya da 10; varsayılan 10)")
    p.add_argument("--dusuk", help="Bu puan ve altı düşük sayılır (varsayılan 10'lukta 6, 5'likte 3)")
    p.add_argument("--min", type=int, default=5, help="Karşılaştırma için en az anket (varsayılan 5)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "csi_analizi.xlsx")
    a = p.parse_args(argv)
    yol = a.anket or ORNEK / "anketler.csv"
    dusuk = para(a.dusuk) if a.dusuk else None
    if (a.dusuk and (dusuk is None or not 0 <= dusuk < a.olcek)) or a.min < 1:
        print("[X] --dusuk 0 ile ölçeğin altı arasında, --min en az 1 olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, a.olcek, dusuk, a.min)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    g = r["genel"]
    print(f"[OK] {g['n']} anket · ortalama {g['ort']}/{a.olcek} · en üst kutu %{g['ust']} · NPS {g['nps']} · düşük puan {g['dusuk']}")
    for s in r["danisman"]:
        print(f"     {s['ad']}: {s['n']} anket · ortalama {s['ort']} · en üst kutu %{s['ust']} · NPS {s['nps']}")
    print(f"     Geri arama listesi: {len(r['dusukler'])} anket ({sum(1 for x in r['dusukler'] if x['geri'])} müşteri aranmak istiyor)")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
