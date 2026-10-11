import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def kural(alan, tur="Metin", brans=None, secenek=(), asgari=None, kosul=None):
    return {"brans": brans, "alan": alan, "anahtar": main.katla(alan), "tur": tur, "secenek": list(secenek), "asgari": D(asgari) if asgari is not None else None,
            "kosul": (kosul[0], main.katla(kosul[0]), main.katla(kosul[1])) if kosul else None, "aciklama": ""}


def talep(no="T1", brans="Yangın", acente="A", **alanlar):
    return {"no": no, "acente": acente, "brans": brans, "sigortali": "", "v": {main.katla(k.replace("_", " ")): v for k, v in alanlar.items()}}


def sonuc(r):
    return [(e["t"]["no"], e["alan"], e["durum"]) for e in r["eksikler"]]


class KuralTesti(unittest.TestCase):
    def test_turler(self):
        k = [kural("Adres", asgari=10), kural("Bedel", "Sayı"), kural("Yıl", "Sayı", asgari=1900), kural("Tarih", "Tarih"), kural("Yapı", "Seçenek", secenek=("Betonarme", "Çelik"))]
        r = main.kontrol([talep(adres="Gül Sokak No: 5", bedel="1.000,50", yil=1900, tarih="01.11.2026", yapi="ÇELİK")], k)
        self.assertEqual((sonuc(r), r["talepler"][0]["sonuc"]), ([], "Tam"))
        r = main.kontrol([talep(adres="Kısa", bedel="0", yil=1899, tarih="yarın", yapi="Ahşap")], k)
        self.assertEqual([e[1:] for e in sonuc(r)], [("Adres", "Geçersiz"), ("Bedel", "Geçersiz"), ("Yıl", "Geçersiz"), ("Tarih", "Geçersiz"), ("Yapı", "Geçersiz")])
        r = main.kontrol([talep(adres="", bedel=None, yil="-", tarih=" ")], k)
        self.assertEqual([e[2] for e in sonuc(r)], ["Eksik"] * 5)
        self.assertEqual((r["talepler"][0]["sonuc"], r["talepler"][0]["eksik"]), ("Eksik var", 5))
        self.assertIn("en az 10", main.deger_kontrol(k[0], "Kısa")[1])

    def test_brans_ve_kosul(self):
        k = [kural("Sigortalı"), kural("Faaliyet", brans={"yangin", "isyeri"}), kural("Alan", "Sayı", brans={"konut"}), kural("Hasar Notu", kosul=("Önceki Hasar", "Var"))]
        r = main.kontrol([talep("T1", "Yangın", sigortali="X", onceki_hasar="Yok"), talep("T2", "KONUT", sigortali="X", alan=90, onceki_hasar="VAR"),
                          talep("T3", "İşyeri", sigortali="X", faaliyet="Depo", onceki_hasar="var", hasar_notu="2025 su baskını")], k)
        self.assertEqual(sonuc(r), [("T1", "Faaliyet", "Eksik"), ("T2", "Hasar Notu", "Eksik")])
        r = main.kontrol([talep("T1", "Nakliyat")], [kural("Faaliyet", brans={"yangin"})])
        self.assertEqual((r["talepler"][0]["sonuc"], [u["tur"] for u in r["uyarilar"]]), ("Kontrol edilmedi", ["Branş için kural yok"]))

    def test_tutarlilik(self):
        k = [kural("Sigortalı")]
        r = main.kontrol([talep("T1", sigortali="X", baslangic_tarihi="01.11.2026", bitis_tarihi="01.11.2026", bina_bedeli="100", muhteviyat_bedeli="50", toplam_sigorta_bedeli="160"),
                          talep("T2", sigortali="X", baslangic_tarihi="01.11.2026", bitis_tarihi="02.11.2026", bina_bedeli="100", muhteviyat_bedeli="", toplam_sigorta_bedeli="100,50"),
                          talep("T3", sigortali="X", toplam_sigorta_bedeli="100")], k)
        self.assertEqual(sonuc(r), [("T1", "Bitiş Tarihi", "Tutarsız"), ("T1", "Toplam Sigorta Bedeli", "Tutarsız")])      # 1 TL'ye kadar fark kabul edilir

    def test_ozet_ve_metin(self):
        k = [kural("İl"), kural("İlçe")]
        r = main.kontrol([talep("T1", acente="A", il="x"), talep("T2", acente="A"), talep("T3", acente="B", il="x", ilce="y")], k, {"il", "ilce"})
        self.assertEqual({a: (v["talep"], v["eksikli"], v["eksik"]) for a, v in r["acente"].items()}, {"A": (2, 2, 3), "B": (1, 0, 0)})
        self.assertEqual((r["alanlar"], r["uyarilar"]), ([("İlçe", 2), ("İl", 1)], []))
        m = main.not_metni(r)
        self.assertIn("=== A ===", m)
        self.assertNotIn("=== B ===", m)
        self.assertIn("  - İlçe: Bilgi girilmemiş.", m)
        self.assertEqual([u["tur"] for u in main.kontrol([talep()], k, {"il"})["uyarilar"]], ["Kural alanı dosyada yok"])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "talepler.csv", ORNEK / "kurallar.csv", cikti)
            d = {t["no"]: t["eksik"] for t in r["talepler"]}
            self.assertEqual(d, {"TT-5001": 0, "TT-5002": 4, "TT-5003": 0, "TT-5004": 1, "TT-5005": 5, "TT-5006": 1, "TT-5007": 0, "TT-5008": 0, "TT-5009": 0, "TT-5010": 2})
            e = {no: sorted((x[1], x[2]) for x in sonuc(r) if x[0] == no) for no in d}
            self.assertEqual(e["TT-5002"], [("Hasar Açıklaması", "Eksik"), ("Risk Adresi", "Geçersiz"), ("Yapı Tarzı", "Geçersiz"), ("İnşa Yılı", "Eksik")])
            self.assertEqual(e["TT-5004"], [("Toplam Sigorta Bedeli", "Tutarsız")])
            self.assertEqual(e["TT-5005"], [("Bitiş Tarihi", "Tutarsız"), ("Brüt Alan (m²)", "Geçersiz"), ("Risk Adresi", "Geçersiz"), ("Önceki Hasar", "Eksik"), ("İlçe", "Eksik")])
            self.assertEqual(e["TT-5010"], [("Hasar Açıklaması", "Geçersiz"), ("İnşa Yılı", "Geçersiz")])
            self.assertEqual({a: (v["talep"], v["eksikli"]) for a, v in r["acente"].items()}, {"Örnek Acente A": (4, 2), "Örnek Acente B": (3, 3), "Örnek Acente C": (3, 0)})
            self.assertEqual(r["uyarilar"], [])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Talep Özeti", "Eksikler", "Acente Özeti", "Alan Özeti", "Uyarılar"])
            self.assertIn("Talep TT-5006", (Path(tmp) / "acente_eksik_notlari.txt").read_text(encoding="utf-8"))

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            t, k = Path(tmp) / "t.csv", Path(tmp) / "k.csv"
            k.write_text("Branş;Alan;Tür;Seçenekler;Asgari;Koşul\n*;İl;;;;\nYangın, İşyeri;Yapı;Seçenek;A|B;;\n*;Kat;Rakam;;;\n*;Tip;Seçenek;;;\n*;Not;Metin;;az;Hasar Var\n*;Yıl;Sayı;;çok;\n",
                         encoding="utf-8")
            kr, uy = main.kurallari_oku(k)
            self.assertEqual([(x["alan"], x["tur"], x["brans"], x["asgari"]) for x in kr], [("İl", "Metin", None, None), ("Yapı", "Seçenek", {"yangin", "isyeri"}, None), ("Yıl", "Sayı", None, None)])
            self.assertEqual(sorted(u["tur"] for u in uy), ["Asgari değer okunamadı", "Asgari değer okunamadı", "Koşul okunamadı", "Kural türü tanınmadı", "Seçenek listesi boş"])
            t.write_text("Talep No;Branş;İl;Yapı;Yıl\nT1;Yangın;Bursa;A;2001\nT1;Yangın;;;\n;Yangın;;;\nT2;Konut;;;0\n", encoding="utf-8")
            tl, sut, uy = main.talepleri_oku(t)
            self.assertEqual(([x["no"] for x in tl], [u["tur"] for u in uy], "yapi" in sut), (["T1", "T2"], ["Mükerrer talep"], True))
            c = str(Path(tmp) / "x.xlsx")
            self.assertEqual(main.main(["--talepler", str(t), "--kurallar", str(k), "--cikti", c]), 0)
            self.assertEqual(main.main(["--talepler", str(Path(tmp) / "yok.csv")]), 1)
            t.write_text("Talep No;Branş\n", encoding="utf-8")
            self.assertEqual(main.main(["--talepler", str(t), "--kurallar", str(k), "--cikti", c]), 1)
            k.write_text("Alan;Tür\nX;Rakam\n", encoding="utf-8")
            self.assertEqual(main.main(["--kurallar", str(k), "--cikti", c]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
