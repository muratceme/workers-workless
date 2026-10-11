"""
Teklif Talebi Bilgi Kontrolü — Workers / Workless kod bloğu
Teknik (Risk Kabul / Underwriting) › Teknik Uzman Yardımcısı

Acenteden gelen teklif taleplerini, sizin verdiğiniz zorunlu alan kurallarına göre kontrol eder.
  Kural   = Branş × Alan × Tür (Metin / Sayı / Tarih / Seçenek) × Seçenekler × Asgari × Koşul
  Eksik   = alan boş            Geçersiz = dolu ama türe, seçeneğe veya asgari değere uymuyor
  Koşul   = "Önceki Hasar=Var" gibi: kural yalnız o alan o değerdeyse uygulanır
  Tutarlılık (sütunlar varsa): bitiş tarihi > başlangıç tarihi; Toplam Sigorta Bedeli = diğer "... Bedeli" sütunlarının toplamı
Zorunlu alan listesi koda gömülü değildir; şirketinizin risk kabul yönetmeliğine göre kural dosyasını siz hazırlarsınız.
Kod bilginin doğruluğunu (adres gerçekten var mı, bedel yeterli mi) denetlemez; yalnız talebin doldurulmuş olmasına bakar.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


KURAL_SUTUN = {"brans": ("brans", "urun"), "alan": ("alan", "bilgi", "sutun"), "tur": ("tur", "tip"), "secenek": ("secenekler", "secenek", "gecerli degerler"),
               "asgari": ("asgari", "en az", "min"), "kosul": ("kosul", "sart"), "aciklama": ("aciklama", "not")}
TURLER = {"metin": "Metin", "sayi": "Sayı", "tarih": "Tarih", "secenek": "Seçenek"}
SABIT = {"no": ("talep no", "teklif no", "talep numarasi"), "acente": ("acente", "acente adi", "acente kodu"), "brans": ("brans", "urun"), "sigortali": ("sigortali", "sigortali adi", "musteri")}


def kurallari_oku(yol: Path) -> tuple[list[dict], list[dict]]:
    kurallar, uy = [], []
    for r in kayitlar(yol, KURAL_SUTUN, ("alan",)):
        alan = metin(r["alan"])
        if not alan:
            continue
        kim = f"{alan} (satır {r['_satir']})"
        tur = TURLER.get(katla(r["tur"]) or "metin")
        if not tur:
            uy.append(uyari("Yüksek", "Kural türü tanınmadı", kim, f"'{metin(r['tur'])}'. Metin, Sayı, Tarih veya Seçenek yazın. Kural uygulanmadı."))
            continue
        secenek = [s.strip() for s in re.split(r"[|;]", metin(r["secenek"])) if s.strip()]
        if tur == "Seçenek" and not secenek:
            uy.append(uyari("Yüksek", "Seçenek listesi boş", kim, "Tür 'Seçenek' ama Seçenekler sütunu boş (ör. Var|Yok). Kural uygulanmadı."))
            continue
        asgari = para(r["asgari"])
        if metin(r["asgari"]) and asgari is None:
            uy.append(uyari("Orta", "Asgari değer okunamadı", kim, f"'{metin(r['asgari'])}'. Asgari değer uygulanmadı."))
        kosul = None
        if metin(r["kosul"]):
            m = re.fullmatch(r"\s*(.+?)\s*=\s*(.+?)\s*", metin(r["kosul"]))
            if not m:
                uy.append(uyari("Yüksek", "Koşul okunamadı", kim, f"'{metin(r['kosul'])}'. 'Alan=Değer' biçiminde yazın. Kural uygulanmadı."))
                continue
            kosul = (m[1], katla(m[1]), katla(m[2]))
        b = [katla(x) for x in re.split(r"[|,;]", metin(r["brans"]))]
        kurallar.append({"brans": None if not any(b) or "*" in metin(r["brans"]) or "hepsi" in b or "tumu" in b else {x for x in b if x}, "alan": alan, "anahtar": katla(alan),
                         "tur": tur, "secenek": secenek, "asgari": asgari, "kosul": kosul, "aciklama": metin(r["aciklama"])})
    if not kurallar:
        raise ValueError(f"{yol.name}: uygulanabilir kural yok.")
    return kurallar, uy


def talepleri_oku(yol: Path) -> tuple[list[dict], set[str], list[dict]]:
    s, uy = tablo_oku(yol), []
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        if any(a in b for a in SABIT["no"]) and any(a in b for a in SABIT["brans"]):
            break
    else:
        raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: Talep No, Branş")
    talepler, gorulen = [], set()
    for n, r in enumerate(s[bi + 1:], bi + 2):
        v = {a: (r[j] if j < len(r) else None) for j, a in enumerate(b) if a}
        sabit = {k: metin(next((v[a] for a in adlar if a in v), "")) for k, adlar in SABIT.items()}
        if not sabit["no"]:
            continue
        if sabit["no"] in gorulen:
            uy.append(uyari("Yüksek", "Mükerrer talep", sabit["no"], f"{yol.name} satır {n}: talep numarası tekrar ediyor; ilk satır kullanıldı."))
            continue
        gorulen.add(sabit["no"])
        talepler.append(sabit | {"v": v, "acente": sabit["acente"] or "Belirtilmemiş"})
    return talepler, {a for a in b if a}, uy


def deger_kontrol(k: dict, ham) -> tuple[str, str] | None:
    """Kurala uymuyorsa (durum, açıklama) döndürür."""
    m = metin(ham)
    if not m or m in ("-", "—", "?"):
        return "Eksik", "Bilgi girilmemiş."
    if k["tur"] == "Metin":
        if k["asgari"] is not None and len(m) < k["asgari"]:
            return "Geçersiz", f"'{m}' çok kısa ({len(m)} karakter; en az {int(k['asgari'])})."
    elif k["tur"] == "Sayı":
        x = para(ham)
        if x is None:
            return "Geçersiz", f"'{m}' sayı olarak okunamadı."
        alt = k["asgari"] if k["asgari"] is not None else None
        if alt is not None and x < alt:
            return "Geçersiz", f"{x:f} asgari değerin ({alt:f}) altında."
        if alt is None and x <= 0:
            return "Geçersiz", f"{x:f} sıfır veya negatif."
    elif k["tur"] == "Tarih":
        if not tarih(ham):
            return "Geçersiz", f"'{m}' tarih olarak okunamadı (GG.AA.YYYY)."
    elif katla(m) not in {katla(s) for s in k["secenek"]}:
        return "Geçersiz", f"'{m}' geçerli seçeneklerden biri değil ({', '.join(k['secenek'])})."
    return None


def kontrol(talepler: list[dict], kurallar: list[dict], sutunlar: set[str] | None = None) -> dict:
    eksikler, uy = [], []
    for t in talepler:
        v, brans, n0 = t["v"], katla(t["brans"]), len(eksikler)
        uygulanan = 0
        for k in kurallar:
            if k["brans"] is not None and brans not in k["brans"]:
                continue
            if k["kosul"] and katla(v.get(k["kosul"][1])) != k["kosul"][2]:
                continue
            uygulanan += 1
            s = deger_kontrol(k, v.get(k["anahtar"]))
            if s:
                eksikler.append({"t": t, "alan": k["alan"], "durum": s[0], "aciklama": s[1], "kural": k["aciklama"]})
        bas, bit = tarih(v.get("baslangic tarihi")), tarih(v.get("bitis tarihi"))
        if bas and bit and bit <= bas:
            eksikler.append({"t": t, "alan": "Bitiş Tarihi", "durum": "Tutarsız", "aciklama": f"Bitiş {bit:%d.%m.%Y}, başlangıçtan ({bas:%d.%m.%Y}) sonra olmalı.", "kural": ""})
        if metin(v.get("toplam sigorta bedeli")):
            toplam, parca = para(v.get("toplam sigorta bedeli")), [para(x) for a, x in v.items() if a.endswith(" bedeli") and a != "toplam sigorta bedeli" and metin(x)]
            if toplam is not None and parca and None not in parca and abs(toplam - sum(parca)) > 1:
                eksikler.append({"t": t, "alan": "Toplam Sigorta Bedeli", "durum": "Tutarsız", "aciklama": f"Toplam {toplam:f}, bedel kalemlerinin toplamı {sum(parca):f}.", "kural": ""})
        if not uygulanan:
            uy.append(uyari("Orta", "Branş için kural yok", t["no"], f"'{t['brans'] or 'boş'}' branşına uyan kural bulunamadı; talep kontrol edilmedi."))
        t["eksik"] = len(eksikler) - n0
        t["sonuc"] = "Kontrol edilmedi" if not uygulanan else "Eksik var" if t["eksik"] else "Tam"
    if sutunlar is not None:
        for alan in sorted({k["alan"] for k in kurallar if k["anahtar"] not in sutunlar}):
            uy.append(uyari("Yüksek", "Kural alanı dosyada yok", alan, "Talep dosyasında bu adda sütun yok; ilgili tüm taleplerde 'Eksik' görünür. Sütun adını veya kuralı düzeltin."))
    acente = defaultdict(lambda: {"talep": 0, "eksikli": 0, "eksik": 0, "alanlar": defaultdict(int)})
    for t in talepler:
        a = acente[t["acente"]]
        a["talep"] += 1
        a["eksikli"] += t["sonuc"] == "Eksik var"
        a["eksik"] += t["eksik"]
    for e in eksikler:
        acente[e["t"]["acente"]]["alanlar"][e["alan"]] += 1
    alanlar = defaultdict(int)
    for e in eksikler:
        alanlar[e["alan"]] += 1
    return {"talepler": talepler, "eksikler": eksikler, "acente": dict(acente), "alanlar": sorted(alanlar.items(), key=lambda x: (-x[1], x[0])), "uyarilar": uy}


def not_metni(r: dict) -> str:
    satir = []
    for ac in sorted(r["acente"]):
        l = [t for t in r["talepler"] if t["acente"] == ac and t["sonuc"] == "Eksik var"]
        if not l:
            continue
        satir += [f"=== {ac} ===", ""]
        for t in l:
            satir.append(f"Talep {t['no']}" + (f" · {t['sigortali']}" if t["sigortali"] else "") + f" · {t['brans']}")
            satir += [f"  - {e['alan']}: {e['aciklama']}" + (f" ({e['kural']})" if e["kural"] else "") for e in r["eksikler"] if e["t"] is t]
            satir.append("")
    return "\n".join(satir) if satir else "Eksiği olan talep yok.\n"


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Talep Özeti"
    _baslik(ws, ["Talep No", "Acente", "Sigortalı", "Branş", "Sonuç", "Eksik / Hatalı Alan", "Alanlar"], (14, 24, 30, 14, 16, 10, 80))
    renk = {"Eksik var": KIRMIZI, "Tam": YESIL, "Kontrol edilmedi": SARI}
    for t in sorted(r["talepler"], key=lambda t: (t["sonuc"] != "Eksik var", -t["eksik"])):
        _satir(ws, [t["no"], t["acente"], t["sigortali"], t["brans"], t["sonuc"], t["eksik"], ", ".join(e["alan"] for e in r["eksikler"] if e["t"] is t)], dolgu={5: renk[t["sonuc"]]}, sar=(7,))
    ws = wb.create_sheet("Eksikler")
    _baslik(ws, ["Talep No", "Acente", "Branş", "Alan", "Durum", "Açıklama", "Kural Notu", "Tamamlandı (siz doldurun)"], (14, 24, 14, 26, 11, 70, 40, 14))
    drenk = {"Eksik": KIRMIZI, "Geçersiz": SARI, "Tutarsız": SARI}
    for e in r["eksikler"]:
        _satir(ws, [e["t"]["no"], e["t"]["acente"], e["t"]["brans"], e["alan"], e["durum"], e["aciklama"], e["kural"], ""], dolgu={5: drenk[e["durum"]]}, sar=(6, 7))
    ws = wb.create_sheet("Acente Özeti")
    _baslik(ws, ["Acente", "Talep", "Eksikli Talep", "Eksikli Talep %", "Toplam Eksik", "En Sık Eksik Alanlar"], (26, 9, 10, 11, 10, 70))
    for ad, a in sorted(r["acente"].items(), key=lambda x: (-x[1]["eksikli"], x[0])):
        sik = sorted(a["alanlar"].items(), key=lambda x: (-x[1], x[0]))[:3]
        _satir(ws, [ad, a["talep"], a["eksikli"], yuvarla(Decimal(a["eksikli"]) / a["talep"] * 100), a["eksik"], ", ".join(f"{k} ({n})" for k, n in sik)], para_s=(4,), sar=(6,))
    ws = wb.create_sheet("Alan Özeti")
    _baslik(ws, ["Alan", "Eksik / Hatalı Talep Sayısı"], (30, 16))
    for alan, n in r["alanlar"]:
        _satir(ws, [alan, n])
    uyari_sayfasi(wb, sirala(r["uyarilar"]))
    kaydet(wb, cikti)
    (cikti.parent / "acente_eksik_notlari.txt").write_text(not_metni(r), encoding="utf-8")


def calistir(talep_yol: Path, kural_yol: Path, cikti: Path) -> dict:
    kurallar, uy1 = kurallari_oku(kural_yol)
    talepler, sutunlar, uy2 = talepleri_oku(talep_yol)
    if not talepler:
        raise ValueError(f"{talep_yol.name}: kontrol edilecek talep yok.")
    r = kontrol(talepler, kurallar, sutunlar)
    r["uyarilar"] = sirala(uy1 + uy2 + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Teklif taleplerinde zorunlu bilgilerin eksiksizliğini kural dosyasına göre kontrol eder.")
    p.add_argument("--talepler", type=Path, default=ORNEK / "talepler.csv", help="Talep No, Acente, Sigortalı, Branş + kurallarda adı geçen sütunlar")
    p.add_argument("--kurallar", type=Path, default=ORNEK / "kurallar.csv", help="Branş, Alan, Tür, Seçenekler, Asgari, Koşul, Açıklama")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "teklif_talebi_kontrol.xlsx")
    a = p.parse_args(argv)
    for y in (a.talepler, a.kurallar):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.talepler, a.kurallar, a.cikti)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    n = len(r["talepler"])
    print(f"[OK] {n} talep · {sum(t['sonuc'] == 'Tam' for t in r['talepler'])} tam · {sum(t['sonuc'] == 'Eksik var' for t in r['talepler'])} eksikli · {len(r['eksikler'])} eksik / hatalı alan")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['kim']} · {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[OK] Acente notları: {(a.cikti.parent / 'acente_eksik_notlari.txt').resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
