import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def o(no, brut, geri=0, tur="A", kol="X", t=date(2026, 1, 15)):
    return {"no": no, "tarih": t, "tur": tur, "kol": kol, "brut": D(str(brut)), "geri": D(str(geri)), "birim": "", "ad": ""}


class AnalizTesti(unittest.TestCase):
    def test_siklik_siddet(self):
        l = [o("1", 1000), o("2", 3000, 1000, kol="Y"), o("3", 10000, tur="a", t=date(2026, 4, 1)), o("4", 200000, 50000, tur="B", t=date(2026, 12, 31)), o("5", 500, 800, tur="B", kol="y", t=date(2025, 3, 31))]
        r = main.analiz(l)
        self.assertEqual({x["no"]: (x["net"], x["dilim"], x["donem"]) for x in r["olaylar"]},
                         {"1": (D(1000), "10.000 altı", "2026 Ç1"), "2": (D(2000), "10.000 altı", "2026 Ç1"), "3": (D(10000), "10.000 – 100.000", "2026 Ç2"), "4": (D(150000), "100.000 – 1.000.000", "2026 Ç4"),
                          "5": (D(-300), "10.000 altı", "2025 Ç1")})
        k = lambda alan: [(s["ad"], s["adet"], s["adet_pay"], s["brut"], s["geri"], s["net"], s["pay"], s["ort"], s["medyan"], s["enb"], s["geri_oran"]) for s in r[alan]]   # noqa: E731
        self.assertEqual(k("turler"), [("B", 2, D("40.00"), D(200500), D(50800), D(149700), D("92.01"), D("74850.00"), D("74850.00"), D(150000), D("25.34")),
                                       ("A", 3, D("60.00"), D(14000), D(1000), D(13000), D("7.99"), D("4333.33"), D(2000), D(10000), D("7.14"))])
        self.assertEqual([(s["ad"], s["adet"], s["net"]) for s in r["kollar"]], [("X", 3, D(161000)), ("Y", 2, D(1700))])
        self.assertEqual([(m["tur"], m["hucre"], m["toplam"]) for m in r["matris"]], [("B", [D(150000), D(-300)], D(149700)), ("A", [D(11000), D(2000)], D(13000))])
        self.assertEqual([(s["ad"], s["adet"], s["net"]) for s in r["dilimler"]], [("10.000 altı", 3, D(2700)), ("10.000 – 100.000", 1, D(10000)), ("100.000 – 1.000.000", 1, D(150000)), ("1.000.000 ve üzeri", 0, D(0))])
        self.assertEqual([(s["ad"], s["adet"], s["net"]) for s in r["donemler"]], [("2025 Ç1", 1, D(-300)), ("2026 Ç1", 2, D(3000)), ("2026 Ç2", 1, D(10000)), ("2026 Ç4", 1, D(150000))])
        oz = r["ozet"]
        self.assertEqual((oz["adet"], oz["brut"], oz["geri"], oz["net"], oz["ort"], oz["medyan"], oz["enb"], oz["enb_pay"], oz["buyuk"], oz["kapsam_disi"], oz["geri_oran"]),
                         (5, D(214500), D(51800), D(162700), D("32540.00"), D(2000), D(150000), D("92.19"), 1, 0, D("24.15")))
        self.assertEqual([x["no"] for x in r["buyukler"]], ["4"])
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]), [("Büyük kayıp olayı", "4"), ("Geri kazanım brüt kayıptan büyük", "5")])

    def test_esik_ve_sinirlar(self):
        l = [o("a", 9999.99), o("b", 10000), o("c", 99999), o("d", 100000), o("e", 1000000), o("f", 50, tur=main.YOK)]
        r = main.analiz(l, esik=D(100), buyuk=D(100000))
        self.assertEqual(({x["no"]: x["dilim"] for x in r["olaylar"]}, r["ozet"]["kapsam_disi"], r["ozet"]["buyuk"], [x["no"] for x in r["buyukler"]]),
                         ({"a": "10.000 altı", "b": "10.000 – 100.000", "c": "10.000 – 100.000", "d": "100.000 – 1.000.000", "e": "1.000.000 ve üzeri"}, 1, 2, ["e", "d"]))
        r = main.analiz([o("a", 60), o("b", 40, tur=main.YOK)], dilimler=(D(50),), buyuk=D(60))
        self.assertEqual(([(s["ad"], s["adet"]) for s in r["dilimler"]], sorted(u["tur"] for u in r["uyarilar"]), r["ozet"]["medyan"]), ([("50 altı", 1), ("50 ve üzeri", 1)], ["Büyük kayıp olayı", "Olay türü yazılmamış"], D("50.00")))
        with self.assertRaises(ValueError):
            main.analiz([o("a", 10)], esik=D(100))


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "kayiplar.csv", cikti)
            oz = r["ozet"]
            self.assertEqual((oz["adet"], oz["brut"], oz["geri"], oz["net"], oz["enb"], oz["buyuk"]), (14, D(2380700), D(199300), D(2181400), D(1200000), 4))
            self.assertEqual([(s["ad"], s["adet"], s["net"]) for s in r["turler"]],
                             [("Müşteriler, ürünler ve iş uygulamaları", 1, D(1200000)), ("İç dolandırıcılık", 1, D(350000)), ("İş kesintisi ve sistem hataları", 2, D(258000)), ("İstihdam uygulamaları ve işyeri güvenliği", 1, D(150000)),
                              ("Dış dolandırıcılık", 3, D(133000)), ("İşlem yürütme, teslimat ve süreç yönetimi", 5, D(75400)), ("Maddi varlıklara zarar", 1, D(15000))])
            self.assertEqual([(s["ad"], s["adet"], s["net"]) for s in r["kollar"]], [("Perakende bankacılık", 8, D(2022100)), ("Ticari bankacılık", 2, D(72800)), ("Alım satım", 1, D(56000)), ("Ödeme ve takas", 3, D(30500))])
            self.assertEqual([(s["ad"], s["adet"]) for s in r["dilimler"]], [("10.000 altı", 3), ("10.000 – 100.000", 7), ("100.000 – 1.000.000", 3), ("1.000.000 ve üzeri", 1)])
            self.assertEqual([(s["ad"], s["adet"]) for s in r["donemler"]], [("2026 Ç1", 5), ("2026 Ç2", 5), ("2026 Ç3", 4)])
            self.assertEqual([x["no"] for x in r["buyukler"]], ["K-007", "K-010", "K-003", "K-006"])
            self.assertEqual([(u["onem"], u["tur"]) for u in r["uyarilar"]], [("Yüksek", "Büyük kayıp olayı")] * 4)
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Olay Türleri", "İş Kolları", "Dönemler", "Tür × İş Kolu", "Tutar Dilimleri", "Büyük Olaylar", "Tüm Olaylar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "k.csv"
            y.write_text("Kayıp No;Tarih;Olay Tipi;Faaliyet Kolu;Kayıp Tutarı;Geri Kazanım Tutarı\nA1;01.03.2026;T;K;1.250,50;250,50\n;02.03.2026;;;500;\nA3;x;T;K;100;\nA4;03.03.2026;T;K;-5;\n;;;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["tarih"], x["tur"], x["kol"], x["brut"], x["geri"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("A1", date(2026, 3, 1), "T", "K", D("1250.50"), D("250.50")), ("Satır 3", date(2026, 3, 2), main.YOK, main.YOK, D(500), D(0))], [("Olay okunamadı", "A3"), ("Olay okunamadı", "A4")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--kayiplar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--esik", "100", "--buyuk", "1000", "--dilimler", "500;1000"]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--esik", "-1"], ["--buyuk", "0"], ["--dilimler", "100;50"], ["--dilimler", "x"], ["--dilimler", ""])], [2] * 5)
            self.assertEqual(main.main(arg + ["--esik", "999999"]), 1)
            self.assertEqual(main.main(["--kayiplar", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Olay No;Açıklama\nA;B\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
