"""
Operasyonel Risk Kayıp Veri Analizi — Gerisi Bende kod bloğu
İç Kontrol ve Risk Yönetimi › Risk Yönetimi Uzmanı

Operasyonel kayıp olaylarını olay türü ve iş kolu bazında toplar; sıklık ve şiddet göstergelerini çıkarır.
  Net kayıp          = brüt kayıp − geri kazanım (sigorta dahil)
  Sıklık             = olay adedi;   sıklık payı % = grubun adedi ÷ toplam adet × 100
  Şiddet             = net kayıp;    ortalama = net ÷ adet;   ortanca ve en büyük olay ayrıca verilir
  Kayıp payı %       = grubun net kaybı ÷ toplam net kayıp × 100
  Geri kazanım %     = geri kazanım ÷ brüt kayıp × 100
  Tutar dilimleri    = brüt kayba göre (--dilimler 10000;100000;1000000; alt sınır dahil, üst sınır hariç)
  Büyük olay         = net kaybı --buyuk (100.000) ve üstünde olan olay
  Kapsam             = brüt kaybı --esik (0) ve üstünde olan olaylar; altındakiler sayılır ama analize girmez
Olay türünü ve iş kolunu siz yazarsınız; kod metni okuyup sınıflandırmaz, sermaye hesabı yapmaz.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"no": ("olay no", "kayip no", "no"), "tarih": ("olay tarihi", "tarih", "tespit tarihi"), "tur": ("olay turu", "kayip olay turu", "olay tipi"), "kol": ("is kolu", "faaliyet kolu"),
         "brut": ("brut kayip", "kayip tutari", "brut kayip tutari"), "geri": ("geri kazanim", "geri kazanim tutari", "sigorta ve diger geri kazanim"), "birim": ("birim", "bolum"), "ad": ("aciklama", "olay ozeti")}
YOK = "(belirtilmemiş)"


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("tarih", "brut")):
        no = metin(r["no"]) or f"Satır {r['_satir']}"
        if not any(metin(r[a]) for a in ("tarih", "brut", "tur", "kol")):
            continue
        t, brut, geri = tarih(r["tarih"]), para(r["brut"]), para(r["geri"]) if metin(r["geri"]) else Decimal(0)
        if t is None or brut is None or brut < 0 or geri is None or geri < 0:
            uy.append(uyari("Yüksek", "Olay okunamadı", no, f"Satır {r['_satir']}: tarih '{metin(r['tarih'])}', brüt kayıp '{metin(r['brut'])}' ya da geri kazanım '{metin(r['geri'])}' okunamadı. Satır alınmadı."))
            continue
        l.append({"no": no, "tarih": t, "tur": metin(r["tur"]) or YOK, "kol": metin(r["kol"]) or YOK, "brut": brut, "geri": geri, "birim": metin(r["birim"]), "ad": metin(r["ad"])})
    return l, uy


def _oran(pay: Decimal, payda: Decimal) -> Decimal | None:
    return yuvarla(Decimal(pay) * 100 / Decimal(payda)) if payda else None


def _medyan(d: list[Decimal]) -> Decimal | None:
    s, n = sorted(d), len(d)
    return None if not n else s[n // 2] if n % 2 else yuvarla((s[n // 2 - 1] + s[n // 2]) / 2)


def _b(x: Decimal) -> str:
    return f"{int(x):,}".replace(",", ".")


def dilim_adlari(dilimler: tuple[Decimal, ...]) -> list[str]:
    return [f"{_b(dilimler[0])} altı"] + [f"{_b(a)} – {_b(u)}" for a, u in zip(dilimler, dilimler[1:])] + [f"{_b(dilimler[-1])} ve üzeri"]


def analiz(olaylar: list[dict], esik: Decimal = Decimal(0), buyuk: Decimal = Decimal(100000), dilimler: tuple[Decimal, ...] = (Decimal(10000), Decimal(100000), Decimal(1000000))) -> dict:
    kapsam = [o for o in olaylar if o["brut"] >= esik]
    if not kapsam:
        raise ValueError("Analiz edilecek kayıp olayı yok.")
    uy, adlar = [], dilim_adlari(dilimler)
    for o in kapsam:
        o["net"] = o["brut"] - o["geri"]
        o["dilim"] = adlar[sum(1 for d in dilimler if o["brut"] >= d)]
        o["donem"] = f"{o['tarih'].year} Ç{(o['tarih'].month - 1) // 3 + 1}"
        if o["geri"] > o["brut"]:
            uy.append(uyari("Orta", "Geri kazanım brüt kayıptan büyük", o["no"], f"Brüt {o['brut']:,.2f}, geri kazanım {o['geri']:,.2f}; net kayıp eksi çıkıyor. Tutarları kontrol edin."))
        if o["net"] >= buyuk:
            uy.append(uyari("Yüksek", "Büyük kayıp olayı", o["no"], f"{o['tarih']:%d.%m.%Y} · {o['tur']} · {o['kol']} · net {o['net']:,.2f}. {o['ad']}"))
        for alan, ad in (("tur", "Olay türü"), ("kol", "İş kolu")):
            if o[alan] == YOK:
                uy.append(uyari("Orta", f"{ad} yazılmamış", o["no"], f"Olay '{YOK}' grubunda toplandı."))
    top_net, top_adet = sum(o["net"] for o in kapsam), len(kapsam)

    def ozet(liste: list[dict]) -> dict:
        brut, geri, net, netler = sum(o["brut"] for o in liste), sum(o["geri"] for o in liste), sum(o["net"] for o in liste), [o["net"] for o in liste]
        return {"adet": len(liste), "brut": brut, "geri": geri, "net": net, "geri_oran": _oran(geri, brut), "ort": yuvarla(net / len(liste)), "medyan": _medyan(netler), "enb": max(netler),
                "adet_pay": _oran(Decimal(len(liste)), Decimal(top_adet)), "pay": _oran(net, top_net) if top_net > 0 else None}

    def kirilim(alan: str, sirali: bool = True) -> list[dict]:
        grup, ilk = {}, []
        for o in kapsam:
            if katla(o[alan]) not in grup:
                ilk.append(katla(o[alan]))
            grup.setdefault(katla(o[alan]), []).append(o)
        s = [{"ad": grup[k][0][alan], **ozet(grup[k])} for k in ilk]
        return sorted(s, key=lambda x: (-x["net"], katla(x["ad"]))) if sirali else s

    turler, kollar = kirilim("tur"), kirilim("kol")
    matris = [{"tur": t["ad"], "hucre": [sum((o["net"] for o in kapsam if katla(o["tur"]) == katla(t["ad"]) and katla(o["kol"]) == katla(k["ad"])), Decimal(0)) for k in kollar], "toplam": t["net"]} for t in turler]
    dil = [{"ad": a, **ozet(g)} if (g := [o for o in kapsam if o["dilim"] == a]) else {"ad": a, "adet": 0, "net": Decimal(0), "brut": Decimal(0), "pay": None, "adet_pay": Decimal("0.00")} for a in adlar]
    donemler = sorted(kirilim("donem", False), key=lambda x: x["ad"])
    o = ozet(kapsam)
    o.update(kapsam_disi=len(olaylar) - len(kapsam), buyuk=sum(1 for x in kapsam if x["net"] >= buyuk), enb_pay=_oran(o["enb"], top_net) if top_net > 0 else None)
    return {"olaylar": sorted(kapsam, key=lambda x: (-x["net"], x["tarih"], x["no"])), "ozet": o, "turler": turler, "kollar": kollar, "matris": matris, "dilimler": dil, "donemler": donemler,
            "buyukler": sorted((x for x in kapsam if x["net"] >= buyuk), key=lambda x: (-x["net"], x["tarih"], x["no"])), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, esik: Decimal, buyuk: Decimal) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    o = r["ozet"]
    _baslik(ws, ["Gösterge", "Değer"], (44, 18))
    for etiket, deger, para_mi in (("Olay adedi (sıklık)", o["adet"], False), ("Brüt kayıp", o["brut"], True), ("Geri kazanım", o["geri"], True), ("Net kayıp", o["net"], True), ("Geri kazanım %", o["geri_oran"], True),
                                   ("Olay başına ortalama net kayıp (şiddet)", o["ort"], True), ("Ortanca net kayıp", o["medyan"], True), ("En büyük olay (net)", o["enb"], True), ("En büyük olayın toplamdaki payı %", o["enb_pay"], True),
                                   (f"Büyük olay adedi (net ≥ {_b(buyuk)})", o["buyuk"], False), (f"Kapsam dışı olay (brüt < {_b(esik)})", o["kapsam_disi"], False)):
        _satir(ws, [etiket, deger], para_s=(2,) if para_mi else ())
    for baslik, alan, ad in (("Olay Türleri", "turler", "Olay Türü"), ("İş Kolları", "kollar", "İş Kolu"), ("Dönemler", "donemler", "Dönem")):
        ws = wb.create_sheet(baslik)
        _baslik(ws, [ad, "Olay Adedi", "Sıklık Payı %", "Brüt Kayıp", "Geri Kazanım", "Net Kayıp", "Kayıp Payı %", "Ortalama", "Ortanca", "En Büyük Olay", "Geri Kazanım %"], (44, 9, 10, 16, 16, 16, 10, 15, 15, 16, 11))
        for s in r[alan]:
            _satir(ws, [s["ad"], s["adet"], s["adet_pay"], s["brut"], s["geri"], s["net"], s["pay"], s["ort"], s["medyan"], s["enb"], s["geri_oran"]], para_s=(3, 4, 5, 6, 7, 8, 9, 10, 11), sar=(1,))
    ws = wb.create_sheet("Tür × İş Kolu")
    _baslik(ws, ["Olay Türü \\ İş Kolu (net kayıp)"] + [k["ad"] for k in r["kollar"]] + ["Toplam"], (44,) + (18,) * (len(r["kollar"]) + 1), "B2")
    for m in r["matris"]:
        _satir(ws, [m["tur"]] + [h or None for h in m["hucre"]] + [m["toplam"]], para_s=range(2, len(m["hucre"]) + 3), sar=(1,))
    ws = wb.create_sheet("Tutar Dilimleri")
    _baslik(ws, ["Brüt Kayıp Dilimi", "Olay Adedi", "Sıklık Payı %", "Net Kayıp", "Kayıp Payı %"], (28, 10, 11, 16, 11))
    for s in r["dilimler"]:
        _satir(ws, [s["ad"], s["adet"], s["adet_pay"], s["net"], s["pay"]], para_s=(3, 4, 5))
    for baslik, liste in (("Büyük Olaylar", r["buyukler"]), ("Tüm Olaylar", r["olaylar"])):
        ws = wb.create_sheet(baslik)
        _baslik(ws, ["Olay No", "Tarih", "Olay Türü", "İş Kolu", "Birim", "Brüt Kayıp", "Geri Kazanım", "Net Kayıp", "Dilim", "Açıklama"], (11, 12, 36, 22, 20, 15, 15, 15, 22, 50))
        for x in liste:
            _satir(ws, [x["no"], x["tarih"], x["tur"], x["kol"], x["birim"], x["brut"], x["geri"], x["net"], x["dilim"], x["ad"]], para_s=(6, 7, 8), tarih_s=(2,), dolgu={8: KIRMIZI if x["net"] >= buyuk else None}, sar=(3, 10))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, esik: Decimal = Decimal(0), buyuk: Decimal = Decimal(100000), dilimler: tuple[Decimal, ...] = (Decimal(10000), Decimal(100000), Decimal(1000000))) -> dict:
    olaylar, uy = oku(yol)
    r = analiz(olaylar, esik, buyuk, dilimler)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, esik, buyuk)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Operasyonel kayıp olaylarını olay türü ve iş kolu bazında toplar; sıklık ve şiddet analizi yapar.")
    p.add_argument("--kayiplar", type=Path, help="Olay No, Olay Tarihi, Olay Türü, İş Kolu, Brüt Kayıp, Geri Kazanım, Birim, Açıklama")
    p.add_argument("--esik", default="0", help="Analize girecek en küçük brüt kayıp (varsayılan 0)")
    p.add_argument("--buyuk", default="100000", help="Büyük olay sayılacak net kayıp (varsayılan 100000)")
    p.add_argument("--dilimler", default="10000;100000;1000000", help="Brüt kayıp dilim sınırları, artan sırada (varsayılan 10000;100000;1000000)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "operasyonel_kayip_analizi.xlsx")
    a = p.parse_args(argv)
    yol = a.kayiplar or ORNEK / "kayiplar.csv"
    esik, buyuk, dilimler = para(a.esik), para(a.buyuk), tuple(para(d) for d in a.dilimler.split(";") if d.strip())
    if esik is None or esik < 0 or buyuk is None or buyuk <= 0 or not dilimler or any(d is None or d <= 0 for d in dilimler) or list(dilimler) != sorted(set(dilimler)):
        print("[X] --esik 0 veya üstü, --buyuk sıfırdan büyük; --dilimler artan sırada pozitif sayılar olmalı (ör. 10000;100000;1000000)")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, esik, buyuk, dilimler)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {o['adet']} olay · brüt {o['brut']:,.2f} · geri kazanım {o['geri']:,.2f} · net {o['net']:,.2f} · büyük olay {o['buyuk']}")
    print(f"[OK] En çok kayıp: {r['turler'][0]['ad']} ({r['turler'][0]['net']:,.2f}) · en sık: {max(r['turler'], key=lambda t: t['adet'])['ad']}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
