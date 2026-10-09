"""
İcra Takip Durum Raporu — Workers / Workless kod bloğu
Hukuk › Avukat

İcra dosyalarının aşama, tahsilat ve masraf bilgilerini tek raporda toplar:
  - Dosya hesabı (rapor tarihine kadar): takip sonrası faiz kalan asıl alacak üzerinden basit faizle (yıllık oran,
    gün/365) işletilir; masraflar eklenir; her tahsilat önce masrafa, sonra faize, sonra asıl alacağa mahsup edilir
    (kısmi ödemenin önce faiz ve giderlere sayılması, TBK 100). Vekâlet ücreti ve harçlar hesaplanmaz.
  - Süre takibi (tebliğ tarihlerinden):
      ödeme emrine itiraz 7 gün (ilamsız İİK 62; rehin İİK 149); kambiyo 5 gün (İİK 168), ödeme süresi 10 gün;
      ilamlı icra emri 7 gün (İİK 32);
      itiraz varsa: itirazın kaldırılması 6 ay (İİK 68), itirazın iptali 1 yıl (İİK 67) — itirazın tebliğinden;
      haciz isteme 1 yıl (İİK 78) — ödeme / icra emrinin tebliğinden (itiraz yoksa hesaplanır).
    Son gün hafta sonu veya ulusal bayrama rastlarsa izleyen iş günü.
  - İşlemsiz dosya: son işlemden bu yana --islemsiz-gun (180) geçmiş açık dosya.
Rapor: özet (aşama dağılımı, alacak / tahsilat / kalan), dosyalar, hesap dökümü, süre takibi, işlemsiz dosyalar,
aylık tahsilat ve masraf, uyarılar. İnternete bağlanmaz.

Kullanım:
    python main.py                                                  # örnek: 12 dosya
    python main.py --dosyalar dosyalar.xlsx --hareketler hareketler.xlsx --bugun 09.10.2026 --islemsiz-gun 180
"""
from __future__ import annotations

import argparse
import calendar
import csv
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
K2 = Decimal("0.01")
SIFIR = Decimal(0)
AYLAR = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]
SABIT_TATILLER = {(1, 1), (4, 23), (5, 1), (5, 19), (7, 15), (8, 30), (10, 29)}
KAPALI = {"kapandi", "kapali", "tahsil edildi", "infaz", "infazen kapandi", "feragat", "sulh", "dustu", "iade", "harcanmis"}
ITIRAZ_VAR = {"var", "evet", "e", "itiraz edildi", "1"}

SUTUNLAR = {"daire": ("icra dairesi", "daire"), "dosya": ("dosya no", "esas no", "dosya"), "borclu": ("borclu",), "tur": ("takip turu", "takip yolu", "tur"),
            "takip_tarihi": ("takip tarihi",), "asil": ("asil alacak", "alacak"), "islemis": ("islemis faiz", "takip oncesi faiz"),
            "oran": ("faiz orani", "yillik faiz", "faiz"), "teblig": ("odeme emri teblig tarihi", "icra emri teblig tarihi", "teblig tarihi"),
            "itiraz": ("itiraz",), "itiraz_teblig": ("itiraz teblig tarihi", "itirazin teblig tarihi"), "haciz": ("haciz tarihi", "ilk haciz tarihi"),
            "asama": ("asama",), "sorumlu": ("sorumlu", "avukat"), "son_islem": ("son islem tarihi", "son islem"), "durum": ("durum",),
            "not": ("not", "aciklama")}
HAREKET_SUTUNLARI = {"dosya": ("dosya no", "esas no", "dosya"), "tarih": ("tarih",), "tur": ("tur", "hareket turu", "islem"), "tutar": ("tutar",),
                     "aciklama": ("aciklama", "not")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("%", "").replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tl(x: Decimal | None) -> str:
    return "—" if x is None else f"{x.quantize(K2, ROUND_HALF_UP):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


def takip_turu(x) -> str:
    k = katla(x)
    if "kambiyo" in k or "cek" in k or "senet" in k or "bono" in k:
        return "Kambiyo"
    if "rehin" in k or "ipotek" in k:
        return "Rehin"
    if "ilamli" in k or ("ilam" in k and "ilamsiz" not in k):
        return "İlamlı"
    return "İlamsız"


def is_gunu(d: date) -> date:
    while d.weekday() >= 5 or (d.month, d.day) in SABIT_TATILLER:
        d += timedelta(days=1)
    return d


def ay_ekle(d: date, n: int) -> date:
    y, m = divmod(d.month - 1 + n, 12)
    yil, ay = d.year + y, m + 1
    return date(yil, ay, min(d.day, calendar.monthrange(yil, ay)[1]))


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Hareket:
    tarih: date
    tur: str          # tahsilat | masraf
    tutar: Decimal
    aciklama: str


@dataclass
class Dosya:
    daire: str
    no: str
    borclu: str
    tur: str
    takip_tarihi: date
    asil: Decimal
    islemis: Decimal
    oran: Decimal | None
    teblig: date | None
    itiraz: bool
    itiraz_teblig: date | None
    haciz: date | None
    asama: str
    sorumlu: str
    son_islem: date | None
    durum: str
    not_: str = ""
    hareketler: list[Hareket] = field(default_factory=list)
    # hesap
    dokum: list[dict] = field(default_factory=list)
    kalan_asil: Decimal = SIFIR
    kalan_faiz: Decimal = SIFIR
    kalan_masraf: Decimal = SIFIR
    toplam_faiz: Decimal = SIFIR
    tahsilat: Decimal = SIFIR
    masraf: Decimal = SIFIR
    sureler: list[dict] = field(default_factory=list)

    @property
    def acik(self) -> bool:
        return katla(self.durum) not in KAPALI

    @property
    def kalan(self) -> Decimal:
        return self.kalan_asil + self.kalan_faiz + self.kalan_masraf

    @property
    def anahtar(self) -> str:
        return katla(f"{self.daire} {self.no}")


def oku(dosya_yolu: Path, hareket_yolu: Path | None) -> tuple[list[Dosya], list[dict]]:
    uy, dosyalar = [], []
    for r in kayitlar(dosya_yolu, SUTUNLAR, ("dosya", "asil")):
        no = metin(r.get("dosya"))
        if not no:
            continue
        tt, asil = tarih(r.get("takip_tarihi")), para(r.get("asil"))
        if not tt or asil is None:
            uy.append({"onem": "Yüksek", "tur": "Eksik dosya bilgisi", "kim": no, "aciklama": "Takip tarihi veya asıl alacak okunamadı; dosya alınmadı"})
            continue
        oran = para(r.get("oran"))
        if oran is not None and oran >= 1:
            oran = oran / 100            # "%24" veya "24" → 0,24
        dosyalar.append(Dosya(metin(r.get("daire")), no, metin(r.get("borclu")), takip_turu(r.get("tur")), tt, asil, para(r.get("islemis")) or SIFIR, oran,
                              tarih(r.get("teblig")), katla(r.get("itiraz")) in ITIRAZ_VAR, tarih(r.get("itiraz_teblig")), tarih(r.get("haciz")),
                              metin(r.get("asama")) or "(belirtilmemiş)", metin(r.get("sorumlu")), tarih(r.get("son_islem")), metin(r.get("durum")) or "Açık",
                              metin(r.get("not"))))
    adlar = Counter(katla(d.no) for d in dosyalar)
    sozluk = {}
    for d in dosyalar:
        sozluk[katla(d.no)] = d if adlar[katla(d.no)] == 1 else None
        sozluk[d.anahtar] = d
    if hareket_yolu:
        for r in kayitlar(hareket_yolu, HAREKET_SUTUNLARI, ("dosya", "tutar")):
            k = katla(r.get("dosya"))
            d = sozluk.get(k)
            t, tutar, tur = tarih(r.get("tarih")), para(r.get("tutar")), katla(r.get("tur"))
            if not k:
                continue
            if d is None:
                uy.append({"onem": "Orta", "tur": "Eşleşmeyen hareket", "kim": metin(r.get("dosya")),
                           "aciklama": f"Satır {r['_satir']}: dosya bulunamadı veya aynı numara birden çok dairede var (Dosya No'ya daire adını ekleyin)"})
                continue
            if not t or tutar is None:
                uy.append({"onem": "Orta", "tur": "Okunamayan hareket", "kim": d.no, "aciklama": f"Satır {r['_satir']}: tarih veya tutar okunamadı"})
                continue
            if "tahsil" in tur or "odeme" in tur or "reddiyat" in tur:
                d.hareketler.append(Hareket(t, "tahsilat", abs(tutar), metin(r.get("aciklama"))))
            elif "masraf" in tur or "harc" in tur or "gider" in tur or "avans" in tur:
                d.hareketler.append(Hareket(t, "masraf", abs(tutar), metin(r.get("aciklama"))))
            else:
                uy.append({"onem": "Orta", "tur": "Bilinmeyen hareket türü", "kim": d.no, "aciklama": f"Satır {r['_satir']}: '{metin(r.get('tur'))}' (Tahsilat / Masraf yazın)"})
    return dosyalar, uy


# ----------------------------------------------------------------------------
# Hesap
# ----------------------------------------------------------------------------

def faiz(tutar: Decimal, oran: Decimal | None, bas: date, son: date) -> Decimal:
    if not oran or son <= bas or tutar <= 0:
        return SIFIR
    return (tutar * oran * (son - bas).days / 365).quantize(K2, ROUND_HALF_UP)


def hesapla(d: Dosya, bugun: date) -> None:
    asil, faiz_b, masraf = d.asil, d.islemis, SIFIR
    d.toplam_faiz = d.islemis
    d.dokum = [{"tarih": d.takip_tarihi, "islem": "Takip", "tutar": d.asil + d.islemis, "faiz_isleyen": SIFIR, "masrafa": SIFIR, "faize": SIFIR, "asla": SIFIR,
                "asil": asil, "faiz": faiz_b, "masraf": masraf, "aciklama": f"Asıl {tl(d.asil)} + işlemiş faiz {tl(d.islemis)}"}]
    onceki = d.takip_tarihi
    for h in sorted((h for h in d.hareketler if h.tarih <= bugun), key=lambda h: (h.tarih, h.tur != "masraf")):
        if h.tarih < d.takip_tarihi:
            continue
        f = faiz(asil, d.oran, onceki, h.tarih)
        faiz_b += f
        d.toplam_faiz += f
        onceki = h.tarih
        satir = {"tarih": h.tarih, "islem": "Masraf" if h.tur == "masraf" else "Tahsilat", "tutar": h.tutar, "faiz_isleyen": f, "masrafa": SIFIR, "faize": SIFIR,
                 "asla": SIFIR, "aciklama": h.aciklama}
        if h.tur == "masraf":
            masraf += h.tutar
            d.masraf += h.tutar
        else:
            d.tahsilat += h.tutar
            kalan = h.tutar
            for alan in ("masrafa", "faize", "asla"):
                bakiye = {"masrafa": masraf, "faize": faiz_b, "asla": asil}[alan]
                m = min(kalan, bakiye)
                satir[alan] = m
                kalan -= m
                if alan == "masrafa":
                    masraf -= m
                elif alan == "faize":
                    faiz_b -= m
                else:
                    asil -= m
            if kalan > 0:
                satir["aciklama"] = (satir["aciklama"] + "; " if satir["aciklama"] else "") + f"fazla tahsilat {tl(kalan)} (borçluya iade / mahsup kontrolü)"
                satir["fazla"] = kalan
        satir.update({"asil": asil, "faiz": faiz_b, "masraf": masraf})
        d.dokum.append(satir)
    f = faiz(asil, d.oran, onceki, bugun)
    faiz_b += f
    d.toplam_faiz += f
    d.dokum.append({"tarih": bugun, "islem": "Rapor tarihi", "tutar": SIFIR, "faiz_isleyen": f, "masrafa": SIFIR, "faize": SIFIR, "asla": SIFIR,
                    "asil": asil, "faiz": faiz_b, "masraf": masraf, "aciklama": "Rapor tarihine kadar işleyen faiz"})
    d.kalan_asil, d.kalan_faiz, d.kalan_masraf = asil, faiz_b, masraf


def sure_ekle(d: Dosya, ad: str, bas: date | None, son: date | None, dayanak: str, bugun: date, durum: str | None = None) -> None:
    if not bas or not son:
        return
    son = is_gunu(son)
    kalan = (son - bugun).days
    d.sureler.append({"ad": ad, "bas": bas, "son": son, "dayanak": dayanak, "kalan": kalan,
                      "durum": durum or ("Süre geçti" if kalan < 0 else "Bugün son gün" if kalan == 0 else f"{kalan} gün kaldı")})


def sureleri_hesapla(d: Dosya, bugun: date) -> None:
    if not d.teblig:
        return
    if d.tur == "Kambiyo":
        sure_ekle(d, "Borçlunun itiraz / şikâyet süresi", d.teblig, d.teblig + timedelta(days=5), "İİK 168", bugun)
        sure_ekle(d, "Ödeme süresi (bitmeden haciz istenemez)", d.teblig, d.teblig + timedelta(days=10), "İİK 168", bugun)
    elif d.tur == "İlamlı":
        sure_ekle(d, "İcra emri: ödeme / itiraz süresi", d.teblig, d.teblig + timedelta(days=7), "İİK 32–33", bugun)
    else:
        sure_ekle(d, "Ödeme emrine itiraz süresi", d.teblig, d.teblig + timedelta(days=7), "İİK 62" if d.tur == "İlamsız" else "İİK 149", bugun)
    if d.itiraz and d.tur in ("İlamsız", "Rehin", "Kambiyo"):
        if d.tur == "Kambiyo":
            return                                       # kambiyoda itiraz icra mahkemesinde incelenir; İİK 67 / 68 süreleri doğmaz
        if d.itiraz_teblig:
            sure_ekle(d, "İtirazın kaldırılması (icra mahkemesi)", d.itiraz_teblig, ay_ekle(d.itiraz_teblig, 6), "İİK 68", bugun)
            sure_ekle(d, "İtirazın iptali davası", d.itiraz_teblig, ay_ekle(d.itiraz_teblig, 12), "İİK 67", bugun)
    elif not d.itiraz:
        durum = f"Haciz yapıldı ({d.haciz:%d.%m.%Y})" if d.haciz else None
        sure_ekle(d, "Haciz isteme süresi", d.teblig, ay_ekle(d.teblig, 12), "İİK 78", bugun, durum)


def analiz_et(dosyalar: list[Dosya], bugun: date, islemsiz_gun: int = 180, uyari_gun: int = 30) -> dict:
    uy = []
    for d in dosyalar:
        hesapla(d, bugun)
        sureleri_hesapla(d, bugun)
        if d.oran is None and d.acik:
            uy.append({"onem": "Orta", "tur": "Faiz oranı yok", "kim": d.no, "aciklama": "Takip sonrası faiz hesaplanmadı; takip talebindeki faiz oranını girin"})
        if any("fazla" in x for x in d.dokum):
            uy.append({"onem": "Orta", "tur": "Fazla tahsilat", "kim": d.no,
                       "aciklama": f"Tahsilat bu hesaptaki borcu {tl(sum((x.get('fazla', SIFIR) for x in d.dokum), SIFIR))} TL aşıyor. Vekâlet ücreti, tahsil "
                                   "harcı ve farklı faiz oranı bu hesapta yoktur; icra dairesi dosya hesabıyla karşılaştırın"})
        if d.acik and d.kalan <= 0:
            uy.append({"onem": "Orta", "tur": "Borç kapanmış, dosya açık", "kim": d.no, "aciklama": "Hesaba göre borç ödenmiş; dosyanın kapatılmasını / infazını isteyin"})
        if not d.acik:
            continue
        if not d.teblig:
            gecen = (bugun - d.takip_tarihi).days
            uy.append({"onem": "Orta" if gecen > 30 else "Bilgi", "tur": "Tebliğ bilgisi yok", "kim": d.no,
                       "aciklama": f"Takip {gecen} gün önce açıldı; ödeme / icra emri tebliğ tarihi yok. Tebligat durumunu sorgulayın"})
        if d.itiraz and d.tur in ("İlamsız", "Rehin") and not d.itiraz_teblig:
            uy.append({"onem": "Orta", "tur": "İtiraz tebliğ tarihi yok", "kim": d.no,
                       "aciklama": "İtiraz var ama itirazın alacaklıya tebliğ tarihi girilmemiş; İİK 67 / 68 süreleri hesaplanamadı"})
        for s in d.sureler:
            if s["ad"].startswith(("Ödeme emrine itiraz", "Borçlunun itiraz", "İcra emri", "Ödeme süresi")):
                continue                                 # borçlu lehine süreler: yalnız bilgi
            if s["durum"].startswith("Haciz yapıldı"):
                continue
            if s["kalan"] < 0:
                uy.append({"onem": "Yüksek", "tur": "Süre geçmiş", "kim": d.no,
                           "aciklama": f"{s['ad']} ({s['dayanak']}): son gün {s['son']:%d.%m.%Y}. İşlem yapıldıysa dosyayı güncelleyin"})
            elif s["kalan"] <= uyari_gun:
                uy.append({"onem": "Yüksek" if s["kalan"] <= 7 else "Orta", "tur": "Yaklaşan süre", "kim": d.no,
                           "aciklama": f"{s['ad']} ({s['dayanak']}): son gün {s['son']:%d.%m.%Y} ({s['kalan']} gün)"})
        if d.son_islem and (bugun - d.son_islem).days > islemsiz_gun:
            uy.append({"onem": "Orta", "tur": "İşlemsiz dosya", "kim": d.no,
                       "aciklama": f"Son işlem {d.son_islem:%d.%m.%Y} ({(bugun - d.son_islem).days} gün önce); haciz / satış / yenileme ihtiyacını değerlendirin"})
    acik = [d for d in dosyalar if d.acik]
    asama = Counter(d.asama for d in acik)
    aylik = defaultdict(lambda: [SIFIR, SIFIR])
    for d in dosyalar:
        for h in d.hareketler:
            if h.tarih <= bugun:
                aylik[(h.tarih.year, h.tarih.month)][0 if h.tur == "tahsilat" else 1] += h.tutar
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    uy.sort(key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))
    return {"dosyalar": dosyalar, "uyarilar": uy, "bugun": bugun, "asama": asama, "aylik": dict(sorted(aylik.items())), "islemsiz_gun": islemsiz_gun}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
UST = Alignment(vertical="top", wrap_text=True)
PF = "#,##0.00"


def _baslik(ws, basliklar, genislik=()):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "B2"


def _f(x):
    return None if x is None else float(x)


def rapor_yaz(cikti: Path, s: dict) -> None:
    dl = s["dosyalar"]
    acik = [d for d in dl if d.acik]
    wb = Workbook()
    oz = wb.active
    oz.title = "Özet"
    _baslik(oz, ["Gösterge", "Değer"], (46, 18))
    alacak = sum((d.asil + d.toplam_faiz + d.masraf for d in dl), SIFIR)
    tahsil = sum((d.tahsilat for d in dl), SIFIR)
    for k, v in [("Rapor tarihi", s["bugun"].strftime("%d.%m.%Y")), ("Dosya sayısı", len(dl)), ("Açık dosya", len(acik)),
                 ("Asıl alacak (tüm dosyalar)", float(sum((d.asil for d in dl), SIFIR))), ("Faiz (işlemiş + takip sonrası)", float(sum((d.toplam_faiz for d in dl), SIFIR))),
                 ("Masraf", float(sum((d.masraf for d in dl), SIFIR))), ("Toplam alacak", float(alacak)), ("Tahsilat", float(tahsil)),
                 ("Tahsilat oranı", float(tahsil / alacak) if alacak else None), ("Kalan alacak (açık dosyalar)", float(sum((d.kalan for d in acik), SIFIR))),
                 ("İşlemsiz açık dosya", sum(u["tur"] == "İşlemsiz dosya" for u in s["uyarilar"]))]:
        oz.append([k, v])
        if isinstance(v, float):
            oz.cell(oz.max_row, 2).number_format = "0.0%" if "oranı" in k else PF
    oz.append([])
    _baslik(oz, ["Aşama (açık dosyalar)", "Dosya", "Kalan Alacak"], ())
    for a, n in s["asama"].most_common():
        oz.append([a, n, float(sum((d.kalan for d in acik if d.asama == a), SIFIR))])
        oz.cell(oz.max_row, 3).number_format = PF
    oz.append([])
    _baslik(oz, ["Takip Türü", "Dosya", "Kalan Alacak"], ())
    for t in ("İlamsız", "Kambiyo", "İlamlı", "Rehin"):
        lst = [d for d in acik if d.tur == t]
        if lst:
            oz.append([t, len(lst), float(sum((d.kalan for d in lst), SIFIR))])
            oz.cell(oz.max_row, 3).number_format = PF

    ds = wb.create_sheet("Dosyalar")
    _baslik(ds, ["İcra Dairesi", "Dosya No", "Borçlu", "Takip Türü", "Takip Tarihi", "Yaş (gün)", "Aşama", "Durum", "Sorumlu", "Asıl Alacak", "Faiz Oranı",
                 "Toplam Faiz", "Masraf", "Tahsilat", "Kalan Asıl", "Kalan Faiz", "Kalan Masraf", "Kalan Toplam", "Tahsilat Oranı", "Son İşlem", "Not"],
            (24, 12, 24, 9, 11, 8, 18, 10, 14, 13, 8, 12, 11, 13, 13, 12, 11, 13, 9, 11, 26))
    for d in sorted(dl, key=lambda d: (not d.acik, -d.kalan)):
        top = d.asil + d.toplam_faiz + d.masraf
        ds.append([d.daire, d.no, d.borclu, d.tur, d.takip_tarihi, (s["bugun"] - d.takip_tarihi).days, d.asama, d.durum, d.sorumlu, float(d.asil), _f(d.oran),
                   float(d.toplam_faiz), float(d.masraf), float(d.tahsilat), float(d.kalan_asil), float(d.kalan_faiz), float(d.kalan_masraf), float(d.kalan),
                   float(d.tahsilat / top) if top else None, d.son_islem, d.not_])
        ds.cell(ds.max_row, 5).number_format = ds.cell(ds.max_row, 20).number_format = "DD.MM.YYYY"
        ds.cell(ds.max_row, 11).number_format = "0.0%"
        ds.cell(ds.max_row, 19).number_format = "0%"
        for j in range(12, 19):
            ds.cell(ds.max_row, j).number_format = PF
        ds.cell(ds.max_row, 10).number_format = PF
    ds.auto_filter.ref = f"A1:U{ds.max_row}"

    hd = wb.create_sheet("Hesap Dökümü")
    _baslik(hd, ["Dosya No", "Tarih", "İşlem", "Tutar", "Dönem Faizi", "Masrafa", "Faize", "Asla", "Kalan Asıl", "Kalan Faiz", "Kalan Masraf", "Kalan Toplam", "Açıklama"],
            (12, 11, 12, 13, 11, 11, 11, 13, 13, 12, 11, 13, 40))
    for d in dl:
        for x in d.dokum:
            hd.append([d.no, x["tarih"], x["islem"], float(x["tutar"]), float(x["faiz_isleyen"]), float(x["masrafa"]), float(x["faize"]), float(x["asla"]),
                       float(x["asil"]), float(x["faiz"]), float(x["masraf"]), float(x["asil"] + x["faiz"] + x["masraf"]), x["aciklama"]])
            hd.cell(hd.max_row, 2).number_format = "DD.MM.YYYY"
            for j in range(4, 13):
                hd.cell(hd.max_row, j).number_format = PF
            if x["islem"] == "Rapor tarihi":
                for j in range(1, 14):
                    hd.cell(hd.max_row, j).font = Font(bold=True)
    hd.auto_filter.ref = f"A1:M{hd.max_row}"

    st = wb.create_sheet("Süre Takibi")
    _baslik(st, ["Dosya No", "Borçlu", "Takip Türü", "Süre", "Başlangıç", "Son Gün", "Dayanak", "Durum", "Yapılan İşlem"], (12, 24, 9, 38, 11, 11, 10, 22, 22))
    for d in dl:
        if not d.acik:
            continue
        for x in d.sureler:
            st.append([d.no, d.borclu, d.tur, x["ad"], x["bas"], x["son"], x["dayanak"], x["durum"], ""])
            st.cell(st.max_row, 5).number_format = st.cell(st.max_row, 6).number_format = "DD.MM.YYYY"
            st.cell(st.max_row, 9).fill = PatternFill("solid", fgColor="FFF4CE")
            if x["kalan"] < 0 and not x["durum"].startswith("Haciz"):
                st.cell(st.max_row, 8).fill = PatternFill("solid", fgColor="FDE2E1")

    isl = wb.create_sheet("İşlemsiz Dosyalar")
    _baslik(isl, ["Dosya No", "Borçlu", "Aşama", "Son İşlem", "Geçen Gün", "Kalan Toplam", "Sorumlu", "Planlanan İşlem"], (12, 24, 18, 11, 9, 14, 14, 26))
    for d in sorted(acik, key=lambda d: d.son_islem or date.min):
        if d.son_islem and (s["bugun"] - d.son_islem).days > s["islemsiz_gun"]:
            isl.append([d.no, d.borclu, d.asama, d.son_islem, (s["bugun"] - d.son_islem).days, float(d.kalan), d.sorumlu, ""])
            isl.cell(isl.max_row, 4).number_format = "DD.MM.YYYY"
            isl.cell(isl.max_row, 6).number_format = PF
            isl.cell(isl.max_row, 8).fill = PatternFill("solid", fgColor="FFF4CE")

    ay = wb.create_sheet("Aylık Tahsilat")
    _baslik(ay, ["Ay", "Tahsilat", "Masraf", "Net"], (14, 14, 14, 14))
    for (y, m), (t, mas) in s["aylik"].items():
        ay.append([f"{AYLAR[m - 1]} {y}", float(t), float(mas), float(t - mas)])
        for j in (2, 3, 4):
            ay.cell(ay.max_row, j).number_format = PF

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Dosya", "Açıklama"], (9, 26, 14, 100))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["kim"], u["aciklama"]])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(dosya_yolu: Path, cikti: Path, bugun: date, hareket_yolu: Path | None = None, **kw) -> dict:
    dosyalar, uy = oku(dosya_yolu, hareket_yolu)
    if not dosyalar:
        raise ValueError(f"{dosya_yolu.name}: icra dosyası bulunamadı")
    s = analiz_et(dosyalar, bugun, **kw)
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    s["uyarilar"] = sorted(uy + s["uyarilar"], key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))
    rapor_yaz(cikti, s)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="İcra dosyalarının aşama, tahsilat, masraf ve sürelerini tek raporda toplar.")
    p.add_argument("--dosyalar", type=Path, default=ORNEK / "dosyalar.csv",
                   help="İcra Dairesi, Dosya No, Borçlu, Takip Türü, Takip Tarihi, Asıl Alacak, İşlemiş Faiz, Faiz Oranı, Ödeme Emri Tebliğ Tarihi, İtiraz, "
                        "İtiraz Tebliğ Tarihi, Haciz Tarihi, Aşama, Sorumlu, Son İşlem Tarihi, Durum")
    p.add_argument("--hareketler", type=Path, help="Dosya No, Tarih, Tür (Tahsilat / Masraf), Tutar, Açıklama")
    p.add_argument("--bugun", help="Rapor tarihi GG.AA.YYYY (örnek veride 09.10.2026)")
    p.add_argument("--islemsiz-gun", type=int, default=180, help="Son işlemden bu kadar gün geçmiş açık dosya işaretlenir (varsayılan 180)")
    p.add_argument("--uyari-gun", type=int, default=30, help="Süre sonuna bu kadar gün kala uyarı (varsayılan 30)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "icra_takip_raporu.xlsx")
    a = p.parse_args(argv)
    ornek = a.dosyalar == ORNEK / "dosyalar.csv"
    har = a.hareketler or (ORNEK / "hareketler.csv" if ornek else None)
    for y in (a.dosyalar, har):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    bugun = tarih(a.bugun) if a.bugun else (date(2026, 10, 9) if ornek else date.today())
    if not bugun:
        print("[X] --bugun GG.AA.YYYY biçiminde olmalı")
        return 2
    try:
        s = calistir(a.dosyalar, a.cikti, bugun, har, islemsiz_gun=a.islemsiz_gun, uyari_gun=a.uyari_gun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    acik = [d for d in s["dosyalar"] if d.acik]
    print(f"[OK] {len(s['dosyalar'])} dosya ({len(acik)} açık) · tahsilat {tl(sum((d.tahsilat for d in s['dosyalar']), SIFIR))} TL · "
          f"açık dosyalarda kalan {tl(sum((d.kalan for d in acik), SIFIR))} TL")
    for u in s["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['kim']} · {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
