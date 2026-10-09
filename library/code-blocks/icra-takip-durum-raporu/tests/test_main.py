import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "i.xlsx"
        cls.s = main.calistir(ORNEK / "dosyalar.csv", cls.cikti, date(2026, 10, 9), ORNEK / "hareketler.csv")
        cls.d = {(d.daire, d.no): d for d in cls.s["dosyalar"]}
        cls.u = {(u["tur"], u["kim"]) for u in cls.s["uyarilar"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_elle_hesaplanan_dosya(self):
        # 250.000 asıl %24: 15.01→10.06 146 gün faiz 24.000; tahsilat 60.000 → masraf 4.850, faiz 18.500 + 24.000, asıl 12.650
        # 10.06→10.08 61 gün: 237.350 × 0,24 × 61 / 365 = 9.520,01; 10.08→09.10 60 gün: 186.870,01 × 0,24 × 60 / 365 = 7.372,41
        d = self.d[("İstanbul 12. İcra Dairesi", "2026/1201 E.")]
        t = [x for x in d.dokum if x["islem"] == "Tahsilat"]
        self.assertEqual((t[0]["faiz_isleyen"], t[0]["masrafa"], t[0]["faize"], t[0]["asla"]), (D("24000.00"), D("4850.00"), D("42500.00"), D("12650.00")))
        self.assertEqual(t[1]["faiz_isleyen"], D("9520.01"))
        self.assertEqual((d.kalan_asil, d.kalan_faiz, d.kalan_masraf), (D("186870.01"), D("7372.41"), D(0)))
        self.assertEqual(d.tahsilat, D(120000))

    def test_fazla_tahsilat_ve_kapanmis_borc(self):
        d = self.d[("İstanbul 7. İcra Dairesi", "2026/455 E.")]
        self.assertEqual(d.dokum[-2]["fazla"], D("160.27"))        # 2.500 + 21.000 + 4.339,73 + 132.000 = 159.839,73
        self.assertEqual(d.kalan, D(0))
        self.assertTrue({("Fazla tahsilat", "2026/455 E."), ("Borç kapanmış, dosya açık", "2026/455 E.")} <= self.u)

    def test_sureler(self):
        s = {(d.no, x["ad"]): x for d in self.s["dosyalar"] for x in d.sureler}
        self.assertEqual(s[("2026/1388 E.", "İtirazın kaldırılması (icra mahkemesi)")]["son"], date(2026, 10, 6))
        self.assertEqual(s[("2026/1388 E.", "İtirazın iptali davası")]["son"], date(2027, 4, 6))
        self.assertEqual(s[("2026/1590 E.", "Borçlunun itiraz / şikâyet süresi")]["son"], date(2026, 10, 12))   # 10.10 Cumartesi
        self.assertEqual(s[("2026/233 E.", "Ödeme emrine itiraz süresi")]["son"], date(2026, 3, 9))            # 08.03 Pazar
        self.assertEqual(s[("2026/610 E.", "Ödeme emrine itiraz süresi")]["dayanak"], "İİK 149")
        ist = self.d[("İstanbul 12. İcra Dairesi", "2026/1201 E.")].sureler
        self.assertTrue(next(x for x in ist if x["ad"] == "Haciz isteme süresi")["durum"].startswith("Haciz yapıldı"))

    def test_uyarilar(self):
        self.assertTrue({("Eksik dosya bilgisi", "2026/1777 E."), ("Süre geçmiş", "2024/7712 E."), ("Süre geçmiş", "2026/1388 E."),
                         ("Yaklaşan süre", "2025/9870 E."), ("İtiraz tebliğ tarihi yok", "2026/610 E."), ("Faiz oranı yok", "2026/233 E."),
                         ("İşlemsiz dosya", "2026/1388 E."), ("Eşleşmeyen hareket", "2026/1201 E."), ("Tebliğ bilgisi yok", "2026/2050 E.")} <= self.u)
        self.assertNotIn(("Yaklaşan süre", "2026/1590 E."), self.u)          # borçlu lehine süreler uyarı üretmez
        self.assertNotIn(("Süre geçmiş", "2025/3321 E."), self.u)            # kapalı dosya

    def test_excel(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Özet", "Dosyalar", "Hesap Dökümü", "Süre Takibi", "İşlemsiz Dosyalar", "Aylık Tahsilat", "Uyarılar"])


class KuralTesti(unittest.TestCase):
    def test_yardimcilar(self):
        self.assertEqual(main.takip_turu("Kambiyo senetlerine özgü (çek)"), "Kambiyo")
        self.assertEqual(main.takip_turu("İpotekli rehin"), "Rehin")
        self.assertEqual(main.takip_turu("İlamlı"), "İlamlı")
        self.assertEqual(main.takip_turu("İlamsız"), "İlamsız")
        self.assertEqual(main.faiz(D(100000), D("0.24"), date(2026, 1, 1), date(2027, 1, 1)), D("24000.00"))
        self.assertEqual(main.ay_ekle(date(2026, 8, 31), 6), date(2027, 2, 28))

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--bugun", "x"]), 2)
            self.assertEqual(main.main(["--dosyalar", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
