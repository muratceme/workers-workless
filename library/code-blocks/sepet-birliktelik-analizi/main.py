"""
Sepet Birliktelik Analizi — Gerisi Bende kod bloğu
CRM ve Sadakat Programı › CRM Uzmanı

Fiş satırlarından birlikte alınan ürün çiftlerini destek, güven ve lift değerleriyle çıkarır.
  Destek %      = çiftin birlikte geçtiği fiş sayısı ÷ toplam fiş sayısı × 100
  Güven A→B %   = çiftin birlikte geçtiği fiş sayısı ÷ A'nın geçtiği fiş sayısı × 100      (A alanların yüzde kaçı B de aldı)
  Lift          = (birlikte fiş × toplam fiş) ÷ (A fiş × B fiş)                            (1'den büyük: beklenenden sık birlikte; 1: ilişki yok; 1'den küçük: beklenenden seyrek)
Bir fişte aynı ürün birden çok satırda geçse de bir kez sayılır; adet ve tutar hesaba girmez. Tek ürünlü fişler toplam fiş sayısına dahildir.
Az sayıda fişe dayanan çiftler yanıltır: --min-fis (varsayılan 5) altındaki çiftler listelenmez. Birliktelik nedensellik değildir.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
from itertools import combinations  # noqa: E402

SUTUN = {"fis": ("fis no", "fis", "fis numarasi", "siparis no", "belge no", "islem no", "sepet no"), "urun": ("urun", "urun adi", "sku", "stok kodu", "urun kodu")}


def oku(yol: Path) -> tuple[dict, dict, list[dict]]:
    """fişler = {fiş no: {ürün anahtarı, ...}}, adlar = {ürün anahtarı: ilk görülen yazım}"""
    fisler, adlar, bos = defaultdict(set), {}, 0
    for r in kayitlar(yol, SUTUN, ("fis", "urun")):
        f, u = metin(r["fis"]), metin(r["urun"])
        if not f or not katla(u):
            bos += bool(f or u)
            continue
        adlar.setdefault(katla(u), u)
        fisler[f].add(katla(u))
    uy = [uyari("Orta", "Eksik satır", yol.name, f"{bos} satırda fiş no veya ürün boş; bu satırlar alınmadı.")] if bos else []
    return dict(fisler), adlar, uy


def analiz(fisler: dict, adlar: dict, min_fis: int = 5, min_destek: Decimal = Decimal(0), min_lift: Decimal = Decimal(0)) -> dict:
    n = len(fisler)
    if not n:
        raise ValueError("Analiz edilecek fiş yok.")
    tek, cift = defaultdict(int), defaultdict(int)
    for urunler in fisler.values():
        for u in urunler:
            tek[u] += 1
        for a, b in combinations(sorted(urunler), 2):
            cift[a, b] += 1
    ciftler = []
    for (a, b), ab in cift.items():
        destek, lift = yuvarla(Decimal(ab) / n * 100), yuvarla(Decimal(ab * n) / (tek[a] * tek[b]))
        if ab < min_fis or destek < min_destek or lift < min_lift:
            continue
        ciftler.append({"a": adlar[a], "b": adlar[b], "birlikte": ab, "a_fis": tek[a], "b_fis": tek[b], "destek": destek, "guven_ab": yuvarla(Decimal(ab) / tek[a] * 100),
                        "guven_ba": yuvarla(Decimal(ab) / tek[b] * 100), "lift": lift, "yorum": "Beklenenden sık birlikte" if lift > 1 else "Beklenenden seyrek birlikte" if lift < 1 else "İlişki yok"})
    return {"ciftler": sorted(ciftler, key=lambda c: (-c["lift"], -c["birlikte"], c["a"], c["b"])),
            "urunler": sorted(({"urun": adlar[u], "fis": k, "destek": yuvarla(Decimal(k) / n * 100)} for u, k in tek.items()), key=lambda x: (-x["fis"], x["urun"])),
            "ozet": {"fis": n, "tek_urunlu": sum(len(u) == 1 for u in fisler.values()), "urun": len(tek), "ortalama": yuvarla(Decimal(sum(len(u) for u in fisler.values())) / n), "cift_toplam": len(cift),
                     "min_fis": min_fis, "min_destek": min_destek, "min_lift": min_lift}}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    o = r["ozet"]
    ws = wb.active
    ws.title = "Ürün Çiftleri"
    _baslik(ws, ["Ürün A", "Ürün B", "Birlikte Fiş", "A Fiş", "B Fiş", "Destek %", "Güven A→B %", "Güven B→A %", "Lift", "Yorum", "Aksiyon (siz doldurun)"], (32, 32, 10, 9, 9, 10, 12, 12, 9, 26, 30))
    for c in r["ciftler"]:
        _satir(ws, [c["a"], c["b"], c["birlikte"], c["a_fis"], c["b_fis"], c["destek"], c["guven_ab"], c["guven_ba"], c["lift"], c["yorum"], ""], para_s=(6, 7, 8, 9),
               dolgu={9: YESIL if c["lift"] > 1 else SARI if c["lift"] < 1 else None})
    ws.append([])
    ws.append([f"{o['fis']} fiş, {o['urun']} ürün. Listelenen: birlikte en az {o['min_fis']} fiş, destek ≥ %{o['min_destek'].normalize():f}, lift ≥ {o['min_lift'].normalize():f}. "
               f"Toplam {o['cift_toplam']} farklı çiftten {len(r['ciftler'])} tanesi koşulları sağladı. Birliktelik nedensellik göstermez."])
    ws = wb.create_sheet("Ürünler")
    _baslik(ws, ["Ürün", "Geçtiği Fiş Sayısı", "Destek %"], (40, 12, 10))
    for u in r["urunler"]:
        _satir(ws, [u["urun"], u["fis"], u["destek"]], para_s=(3,))
    ws = wb.create_sheet("Özet")
    _baslik(ws, ["Gösterge", "Değer"], (40, 14))
    for ad, d in (("Fiş sayısı", o["fis"]), ("Tek ürünlü fiş", o["tek_urunlu"]), ("Ürün çeşidi", o["urun"]), ("Fiş başına ortalama ürün çeşidi", o["ortalama"]), ("Görülen farklı çift", o["cift_toplam"]),
                  ("Listelenen çift", len(r["ciftler"]))):
        _satir(ws, [ad, d])
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, min_fis: int = 5, min_destek: Decimal = Decimal(0), min_lift: Decimal = Decimal(0)) -> dict:
    fisler, adlar, uy = oku(yol)
    r = analiz(fisler, adlar, min_fis, min_destek, min_lift)
    if r["ozet"]["fis"] < 100:
        uy.append(uyari("Orta", "Fiş sayısı az", yol.name, f"{r['ozet']['fis']} fiş var. Az veriyle bulunan çiftler rastlantı olabilir; daha uzun dönemle çalıştırın."))
    if not r["ciftler"]:
        uy.append(uyari("Bilgi", "Koşulları sağlayan çift yok", yol.name, "Eşikleri (--min-fis, --min-destek, --min-lift) düşürmeyi veya daha çok fişle çalıştırmayı deneyin."))
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Fiş satırlarından birlikte alınan ürün çiftlerini destek, güven ve lift ile çıkarır.")
    p.add_argument("--fisler", type=Path, default=ORNEK / "fis_satirlari.csv", help="Fiş No, Ürün (her satır bir fiş kalemi)")
    p.add_argument("--min-fis", help="Çiftin birlikte geçmesi gereken en az fiş sayısı (varsayılan 5; örnek veride 2)")
    p.add_argument("--min-destek", default="0", help="En az destek %% (varsayılan 0)")
    p.add_argument("--min-lift", default="0", help="En az lift (varsayılan 0: hepsi; yalnız olumlu ilişki için 1)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "sepet_birliktelik.xlsx")
    a = p.parse_args(argv)
    mf, md, ml = para(a.min_fis or ("2" if a.fisler == ORNEK / "fis_satirlari.csv" else "5")), para(a.min_destek), para(a.min_lift)
    if mf is None or mf < 1 or mf != int(mf) or md is None or ml is None or md < 0 or ml < 0:
        print("[X] --min-fis 1 veya daha büyük tam sayı; --min-destek ve --min-lift sıfır veya pozitif sayı olmalı")
        return 2
    if not a.fisler.exists():
        print(f"[X] Dosya bulunamadı: {a.fisler}")
        return 1
    try:
        r = calistir(a.fisler, a.cikti, int(mf), md, ml)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {o['fis']} fiş · {o['urun']} ürün · {o['cift_toplam']} farklı çift, {len(r['ciftler'])} tanesi listelendi")
    for c in r["ciftler"][:5]:
        print(f"     {c['a']} + {c['b']}: lift {c['lift']} · destek %{c['destek']} · güven %{c['guven_ab']} / %{c['guven_ba']} ({c['birlikte']} fiş)")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
