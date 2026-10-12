import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def fis(*sepetler):
    f = {str(i): {x.lower() for x in s} for i, s in enumerate(sepetler)}
    return f, {x: x.upper() for s in f.values() for x in s}


class AnalizTesti(unittest.TestCase):
    def test_destek_guven_lift(self):
        r = main.analiz(*fis("ab", "ab", "ab", "a", "a", "b", "c", "c", "c", "c"), min_fis=1)
        c = r["ciftler"][0]
        self.assertEqual((c["a"], c["b"], c["birlikte"], c["a_fis"], c["b_fis"], c["destek"], c["guven_ab"], c["guven_ba"], c["lift"], c["yorum"]),
                         ("A", "B", 3, 5, 4, D("30.00"), D("60.00"), D("75.00"), D("1.50"), "Beklenenden sık birlikte"))
        self.assertEqual(main.analiz(*fis("ab", "a", "b", "c"), min_fis=1)["ciftler"][0]["lift"], D("1.00"))           # 1 × 4 ÷ (2 × 2)
        self.assertEqual(main.analiz(*fis("ab", "a", "a", "b", "b"), min_fis=1)["ciftler"][0]["yorum"], "Beklenenden seyrek birlikte")   # 1 × 5 ÷ 9 = 0,56
        self.assertEqual(r["ozet"] | {"min_destek": 0, "min_lift": 0}, {"fis": 10, "tek_urunlu": 7, "urun": 3, "ortalama": D("1.30"), "cift_toplam": 1, "min_fis": 1, "min_destek": 0, "min_lift": 0})
        self.assertEqual([(u["urun"], u["fis"], u["destek"]) for u in r["urunler"]], [("A", 5, D("50.00")), ("B", 4, D("40.00")), ("C", 4, D("40.00"))])
        with self.assertRaises(ValueError):
            main.analiz({}, {})

    def test_esikler_ve_siralama(self):
        f = fis("ab", "ab", "ab", "ac", "cd", "cd", "d", "e", "e", "e")
        self.assertEqual([(c["a"], c["b"], c["birlikte"], c["lift"]) for c in main.analiz(*f, min_fis=1)["ciftler"]], [("A", "B", 3, D("2.50")), ("C", "D", 2, D("2.22")), ("A", "C", 1, D("0.83"))])
        self.assertEqual(len(main.analiz(*f, min_fis=2)["ciftler"]), 2)
        self.assertEqual(len(main.analiz(*f, min_fis=1, min_destek=D(25))["ciftler"]), 1)
        self.assertEqual(len(main.analiz(*f, min_fis=1, min_lift=D(1))["ciftler"]), 2)
        self.assertEqual(main.analiz(*f)["ciftler"], [])                                                                # varsayılan en az 5 fiş
        self.assertEqual(len(main.analiz(*fis("abc"), min_fis=1)["ciftler"]), 3)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "s.xlsx"
            r = main.calistir(ORNEK / "fis_satirlari.csv", cikti, min_fis=2)
            self.assertEqual([(c["a"].split()[0], c["b"].split()[0], c["birlikte"], c["destek"], c["guven_ab"], c["guven_ba"], c["lift"]) for c in r["ciftler"]],
                             [("Filtre", "Süt", 2, D("20.00"), D("66.67"), D("66.67"), D("2.22")), ("Çay", "Toz", 3, D("30.00"), D("60.00"), D("75.00"), D("1.50")),
                              ("Bisküvi", "Çay", 2, D("20.00"), D("66.67"), D("40.00"), D("1.33"))])
            self.assertEqual((r["ozet"]["fis"], r["ozet"]["tek_urunlu"], r["ozet"]["urun"], r["ozet"]["cift_toplam"], [u["tur"] for u in r["uyarilar"]]), (10, 3, 6, 7, ["Fiş sayısı az"]))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Ürün Çiftleri", "Ürünler", "Özet", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "f.csv"
            y.write_text("Sipariş No;Ürün Adı\n1;Çay\n1;ÇAY\n1;Şeker\n2;çay\n;Şeker\n3;\n", encoding="utf-8")
            f, adlar, uy = main.oku(y)
            self.assertEqual((f, adlar, [u["tur"] for u in uy]), ({"1": {"cay", "seker"}, "2": {"cay"}}, {"cay": "Çay", "seker": "Şeker"}, ["Eksik satır"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--fisler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--min-fis", "1", "--min-destek", "10", "--min-lift", "0,5"]), 0)
            self.assertEqual(main.main(arg), 0)                                         # koşulu sağlayan çift yok: yine rapor üretir
            self.assertEqual(main.main(arg + ["--min-fis", "0"]), 2)
            self.assertEqual(main.main(arg + ["--min-fis", "2,5"]), 2)
            self.assertEqual(main.main(["--fisler", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Fiş No;Ürün\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
