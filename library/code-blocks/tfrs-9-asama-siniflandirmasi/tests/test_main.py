import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
OLCEK = ["AAA", "AA", "A", "BBB", "BB", "B", "CCC", "CC", "C"]


def kredi(no="K1", gun=0, bakiye=1000, musteri=None, segment="Bireysel", onceki=None, **kw):
    k = {"no": no, "musteri_no": musteri or no, "musteri": "M", "segment": segment, "urun": "", "bakiye": D(str(bakiye)), "gun": gun, "yy": False, "yi": False, "temerrut": False,
         "ilk_not": None, "not": None, "ilk_not_ad": "", "not_ad": "", "ilk_to": None, "to": None, "onceki_asama": onceki}
    for a, v in kw.items():
        k[a] = D(str(v)) if a in ("ilk_to", "to") else v
    return k


class KuralTesti(unittest.TestCase):
    def test_gecikme_esikleri(self):
        self.assertEqual([main.siniflandir(kredi(gun=g))[0] for g in (0, 30, 31, 90, 91)], [1, 1, 2, 2, 3])
        self.assertEqual(main.siniflandir(kredi(gun=45), a2_gun=60)[0], 1)
        self.assertEqual(main.siniflandir(kredi(gun=95), a3_gun=180)[0], 2)
        self.assertEqual(main.siniflandir(kredi(gun=45))[1], ["Gecikme 45 gün (> 30)"])

    def test_nitel_olcutler(self):
        self.assertEqual(main.siniflandir(kredi(yy=True))[0], 2)
        self.assertEqual(main.siniflandir(kredi(yi=True))[0], 2)
        self.assertEqual(main.siniflandir(kredi(temerrut=True, yy=True)), (3, ["Temerrüt işareti", "Yeniden yapılandırma"]))

    def test_not_ve_to(self):
        def n(ilk, guncel, **kw):
            return main.siniflandir(kredi(ilk_not=OLCEK.index(ilk), ilk_not_ad=ilk, **{"not": OLCEK.index(guncel), "not_ad": guncel}), **kw)
        self.assertEqual(n("A", "BB")[0], 1)                                          # 2 kademe
        self.assertEqual(n("A", "B"), (2, ["Not 3 kademe kötüleşti (A → B)"]))
        self.assertEqual(n("A", "BB", not_esik=2)[0], 2)
        self.assertEqual(n("B", "A")[0], 1)                                           # iyileşme
        self.assertEqual(main.siniflandir(kredi(ilk_to="1.0", to="1.99"))[0], 1)
        self.assertEqual(main.siniflandir(kredi(ilk_to="1.0", to="2.0"))[0], 2)
        self.assertEqual(main.siniflandir(kredi(ilk_to="1.0", to="2.5"), to_kat=D(3))[0], 1)
        self.assertEqual(main.not_sirasi("bbb", OLCEK), 3)
        self.assertEqual(main.not_sirasi("7", []), 7)
        self.assertIsNone(main.not_sirasi("D", OLCEK))

    def test_musteri_bazi(self):
        k = [kredi("T1", 100, musteri="M1", segment="Ticari"), kredi("T2", 0, musteri="M1", segment="Ticari"), kredi("B1", 100, musteri="M2"), kredi("B2", 0, musteri="M2")]
        r = main.analiz(k, bulasma={"ticari"})
        self.assertEqual({x["no"]: (x["asama"], x["bulasma"]) for x in r["krediler"]}, {"T1": (3, False), "T2": (3, True), "B1": (3, False), "B2": (1, False)})
        self.assertEqual(next(x for x in k if x["no"] == "T2")["gerekce"], ["Müşteri bazı: T1 Aşama 3"])
        self.assertEqual([(u["tur"], u["kim"]) for u in r["uyarilar"]], [("Müşterinin kredileri farklı aşamalarda", "M2")])

    def test_ozet(self):
        r = main.analiz([kredi("A", 0, 600), kredi("B", 40, 300, segment="KOBİ"), kredi("C", 200, 100, segment="KOBİ")])
        self.assertEqual({a: (v["adet"], v["bakiye"], v["pay"]) for a, v in r["ozet"].items()}, {1: (1, D(600), D("60.00")), 2: (1, D(300), D("30.00")), 3: (1, D(100), D("10.00"))})
        self.assertEqual((r["segment"]["KOBİ"]["a2_pay"], r["segment"]["KOBİ"]["a3_pay"]), (D("75.00"), D("25.00")))
        self.assertEqual(r["gerekce"], {(2, "Gecikme"): [1, D(300)], (3, "Gecikme"): [1, D(100)]})
        self.assertIsNone(r["gecis"])


class GecisTesti(unittest.TestCase):
    def test_onceki_dosya(self):
        eski = [kredi("A", 0, 100), kredi("B", 0, 200), kredi("C", 50, 300), kredi("D", 120, 400), kredi("E", 0, 500, onceki=2)]
        yeni = [kredi("A", 45, 90), kredi("B", 0, 190), kredi("C", 0, 280), kredi("D", 150, 390), kredi("F", 95, 50)]
        r = main.analiz(yeni, eski)
        g = r["gecis"]
        self.assertEqual((g[1][2], g[1][1], g[2][1], g[3][3], g[2]["Kapandı"], g["Yeni"][3]), (D(100), D(200), D(300), D(400), D(500), D(50)))   # E: Aşama sütunu 2
        self.assertEqual([(d["k"]["no"], d["onceki"], d["k"]["asama"], d["yon"]) for d in r["degisen"]], [("A", 1, 2, "Kötüleşti"), ("C", 2, 1, "İyileşti")])
        self.assertEqual([u["tur"] for u in r["uyarilar"]], ["Aşama 2 → 1 iyileşme"])

    def test_ayni_dosyada_onceki_asama(self):
        r = main.analiz([kredi("A", 95, onceki=1), kredi("B", 0, onceki=1), kredi("C", 0)])
        self.assertEqual([(d["k"]["no"], d["yon"]) for d in r["degisen"]], [("A", "Kötüleşti")])
        self.assertEqual(r["gecis"]["Yeni"][1], D(1000))


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "krediler.csv", cikti, ORNEK / "onceki_donem.csv", OLCEK, {"ticari", "kobi"})
            self.assertEqual([r["ozet"][a]["adet"] for a in (1, 2, 3)], [30, 9, 7])
            self.assertEqual(sum(v["bakiye"] for v in r["ozet"].values()), r["toplam"])
            d = {x["k"]["no"]: (x["onceki"], x["k"]["asama"]) for x in r["degisen"]}
            self.assertEqual((d["KR910129"], d["KR910105"], d["KR910102"]), ((1, 3), (1, 2), (1, 2)))        # müşteri bazı · not · TO
            self.assertTrue(next(k for k in r["krediler"] if k["no"] == "KR910129")["bulasma"])
            self.assertEqual([u["kim"] for u in r["uyarilar"] if u["tur"].startswith("Müşterinin")], ["M40110"])   # bireysel: bulaşma yok, uyarı var
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Sınıflandırma", "Aşama Değişenler", "Geçiş Matrisi", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "k.csv"
            y.write_text("Kredi No;Risk Bakiyesi;Gecikme Gün;Yeniden Yapılandırma;Temerrüt;İlk Not;Güncel Not;İlk TO (%);Güncel TO (%)\n"
                         "K1;1.000,00;35;;;A;BBB;0,5;0,6\nK2;2.000,00;;Evet;;A;ZZ;;\nK3;500,00;çok;;X;;;;\nK1;1;0;;;;;;\nK4;eksi;0;;;;;;\n", encoding="utf-8")
            k, uy = main.oku(y, OLCEK)
            self.assertEqual([(x["no"], x["gun"], x["yy"], x["temerrut"], x["ilk_not"], x["not"], x["to"]) for x in k],
                             [("K1", 35, False, False, 2, 3, D("0.6")), ("K2", 0, True, False, 2, None, None), ("K3", 0, False, True, None, None, None)])
            self.assertEqual(sorted(u["tur"] for u in uy), ["Gecikme günü okunamadı", "Not ölçekte yok", "Satır okunamadı", "Satır okunamadı"])
            arg = ["--krediler", str(y), "--not-olcegi", ",".join(OLCEK), "--cikti", str(Path(tmp) / "x.xlsx")]
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(load_workbook(Path(tmp) / "x.xlsx").sheetnames, ["Özet", "Sınıflandırma", "Uyarılar"])
            self.assertEqual(main.main(arg + ["--asama2-gun", "90"]), 2)
            self.assertEqual(main.main(["--krediler", str(Path(tmp) / "yok.csv")]), 1)
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "o.xlsx")]), 0)


if __name__ == "__main__":
    unittest.main()
