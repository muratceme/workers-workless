"""
TFRS 9 Aşama Sınıflandırması — Workers / Workless kod bloğu
Krediler İzleme ve Takip › Kredi İzleme Uzmanı

Kredileri gecikme günü ve risk göstergelerine göre Aşama 1-2-3 olarak sınıflandırır; her kredinin gerekçesini yazar ve
önceki döneme göre aşaması değişenleri listeler.
  Aşama 3 (temerrüt / değer düşüklüğü):  gecikme > --temerrut-gun (varsayılan 90)  veya  "Temerrüt" işareti.
  Aşama 2 (kredi riskinde önemli artış): gecikme > --asama2-gun (varsayılan 30)  veya  yeniden yapılandırma
                                         veya  yakın izleme işareti
                                         veya  not kötüleşmesi ≥ --not-esik kademe (ilk not → güncel not)
                                         veya  güncel TO ÷ ilk TO ≥ --to-kat  (TO = temerrüt olasılığı)
  Aşama 1: diğerleri.
  Gerekçeler birikir (bir kredi birden çok ölçütle Aşama 2 olabilir); en yüksek aşama geçerlidir.
  Müşteri bazı (--bulasma "Ticari,KOBİ"): bu segmentlerde bir kredisi Aşama 3 olan müşterinin tüm kredileri Aşama 3 olur.
  Not ölçeği --not-olcegi ile en iyiden en kötüye verilir (ör. "AAA,AA,A,BBB,BB,B,CCC"); sayısal notta küçük sayı iyidir.
TFRS 9'da 30 günü aşan gecikme kredi riskinde önemli artışın, 90 günü aşan gecikme temerrüdün aksi ispat edilebilir
karinesidir; eşikler ve nitel ölçütler bankanızın politikasına göre belirlenir. Beklenen kredi zararı (karşılık)
hesaplanmaz. İnternete bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: 48 kredi, önceki dönemle
    python main.py --krediler k.xlsx --onceki gecen_ay.xlsx --bulasma "Ticari,KOBİ" --not-olcegi "AAA,AA,A,BBB,BB,B,CCC,CC,C"
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUNLAR = {"no": ("kredi no", "kredi numarasi", "hesap no"), "musteri_no": ("musteri no", "musteri numarasi"), "musteri": ("musteri", "musteri adi", "unvan"),
            "segment": ("segment", "portfoy"), "urun": ("urun", "kredi turu"), "bakiye": ("risk bakiyesi", "bakiye", "risk tutari"),
            "gun": ("gecikme gun", "gecikme gunu", "gecikme"), "yy": ("yeniden yapilandirma", "yapilandirma", "yeniden yapilandirildi"),
            "yi": ("yakin izleme", "izleme listesi", "yakin izlemede"), "temerrut": ("temerrut", "donuk", "tga", "takip"),
            "ilk_not": ("ilk not", "acilis notu", "ilk derecelendirme notu"), "not": ("guncel not", "not", "derecelendirme notu"),
            "ilk_to": ("ilk to", "ilk pd", "acilis pd", "ilk temerrut olasiligi"), "to": ("guncel to", "guncel pd", "pd", "temerrut olasiligi"),
            "asama": ("asama", "onceki asama")}
EVET = ("evet", "e", "var", "1", "x", "true", "yes")


def evet_mi(x) -> bool:
    return katla(x) in EVET


def not_sirasi(x, olcek: list[str]) -> int | None:
    """Ölçekteki sıra (0 = en iyi). Ölçek yoksa sayısal not kullanılır."""
    s = metin(x).upper()
    if not s:
        return None
    if olcek:
        return olcek.index(s) if s in olcek else None
    v = para(s)
    return int(v) if v is not None and v == v.to_integral_value() else None


def oku(yol: Path, olcek: list[str]) -> tuple[list[dict], list[dict]]:
    uy, krediler, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUNLAR, ("no", "bakiye")):
        no = metin(r["no"])
        if not no:
            continue
        b = para(r["bakiye"])
        if no in gorulen or b is None or b < 0:
            uy.append(uyari("Yüksek", "Satır okunamadı", no, f"{yol.name} satır {r['_satir']}: " + ("kredi numarası tekrar ediyor" if no in gorulen else f"bakiye '{metin(r['bakiye'])}'") +
                            ". Satır atlandı."))
            continue
        gorulen.add(no)
        gun = para(r["gun"])
        if metin(r["gun"]) and gun is None:
            uy.append(uyari("Yüksek", "Gecikme günü okunamadı", no, f"'{metin(r['gun'])}'; 0 alındı."))
        notlar = {}
        for alan in ("ilk_not", "not"):
            notlar[alan] = not_sirasi(r[alan], olcek)
            if metin(r[alan]) and notlar[alan] is None:
                uy.append(uyari("Orta", "Not ölçekte yok", no, f"'{metin(r[alan])}' not ölçeğinde tanımlı değil; not ölçütü bu kredi için uygulanmadı."))
        asama = para(r["asama"])
        krediler.append({"no": no, "musteri_no": metin(r["musteri_no"]) or no, "musteri": metin(r["musteri"]), "segment": metin(r["segment"]) or "Belirtilmemiş",
                         "urun": metin(r["urun"]), "bakiye": b, "gun": max(int(gun or 0), 0), "yy": evet_mi(r["yy"]), "yi": evet_mi(r["yi"]), "temerrut": evet_mi(r["temerrut"]),
                         "ilk_not": notlar["ilk_not"], "not": notlar["not"], "ilk_not_ad": metin(r["ilk_not"]), "not_ad": metin(r["not"]),
                         "ilk_to": para(r["ilk_to"]), "to": para(r["to"]), "onceki_asama": int(asama) if asama in (1, 2, 3) else None})
    return krediler, uy


def siniflandir(k: dict, a2_gun: int = 30, a3_gun: int = 90, not_esik: int = 3, to_kat: Decimal = Decimal(2)) -> tuple[int, list[str]]:
    g3, g2 = [], []
    if k["gun"] > a3_gun:
        g3.append(f"Gecikme {k['gun']} gün (> {a3_gun})")
    if k["temerrut"]:
        g3.append("Temerrüt işareti")
    if a2_gun < k["gun"] <= a3_gun:
        g2.append(f"Gecikme {k['gun']} gün (> {a2_gun})")
    if k["yy"]:
        g2.append("Yeniden yapılandırma")
    if k["yi"]:
        g2.append("Yakın izleme")
    if k["ilk_not"] is not None and k["not"] is not None and k["not"] - k["ilk_not"] >= not_esik:
        g2.append(f"Not {k['not'] - k['ilk_not']} kademe kötüleşti ({k['ilk_not_ad']} → {k['not_ad']})")
    if k["ilk_to"] and k["to"] and k["ilk_to"] > 0 and k["to"] / k["ilk_to"] >= to_kat:
        g2.append(f"TO {yuvarla(k['to'] / k['ilk_to'], '0.1')} katına çıktı (%{k['ilk_to']:f} → %{k['to']:f})")
    if g3:
        return 3, g3 + g2
    return (2, g2) if g2 else (1, [])


def analiz(krediler: list[dict], onceki: list[dict] | None = None, bulasma: set[str] = frozenset(), **esik) -> dict:
    uy = []
    for k in krediler:
        k["asama"], k["gerekce"] = siniflandir(k, **esik)
        k["bulasma"] = False
    musteri = defaultdict(list)
    for k in krediler:
        musteri[k["musteri_no"]].append(k)
    karisik = []
    for no, l in musteri.items():
        if len({k["asama"] for k in l}) > 1 and any(k["asama"] == 3 for k in l):
            kaynak = next(k for k in l if k["asama"] == 3)
            for k in l:
                if k["asama"] < 3 and katla(k["segment"]) in bulasma:
                    k["asama"], k["bulasma"] = 3, True
                    k["gerekce"] = [f"Müşteri bazı: {kaynak['no']} Aşama 3"] + k["gerekce"]
            if len({k["asama"] for k in l}) > 1:
                karisik.append((no, l))
    for no, l in karisik:
        uy.append(uyari("Orta", "Müşterinin kredileri farklı aşamalarda", no, "Aşama 3 kredisi olan müşterinin diğer kredileri: " +
                        ", ".join(f"{k['no']} (Aşama {k['asama']}, {k['segment']})" for k in l if k["asama"] < 3) +
                        ". Borçlu bazında değerlendirme gerekip gerekmediğini politikanıza göre kontrol edin."))
    for k in krediler:
        if k["temerrut"] and k["gun"] == 0:
            uy.append(uyari("Bilgi", "Temerrüt işareti var, gecikme yok", k["no"], "Gecikmesiz temerrüt (ödeme güçlüğü göstergesi) olabilir; kaydı doğrulayın."))
    toplam = sum((k["bakiye"] for k in krediler), Decimal(0))
    ozet = {a: {"adet": sum(k["asama"] == a for k in krediler), "bakiye": sum((k["bakiye"] for k in krediler if k["asama"] == a), Decimal(0))} for a in (1, 2, 3)}
    for a in ozet:
        ozet[a]["pay"] = yuvarla(ozet[a]["bakiye"] / toplam * 100) if toplam else Decimal(0)
    segment = {}
    for s in sorted({k["segment"] for k in krediler}):
        l = [k for k in krediler if k["segment"] == s]
        st = sum((k["bakiye"] for k in l), Decimal(0))
        segment[s] = {"adet": len(l), "bakiye": st} | {a: sum((k["bakiye"] for k in l if k["asama"] == a), Decimal(0)) for a in (1, 2, 3)}
        segment[s]["a2_pay"], segment[s]["a3_pay"] = (yuvarla(segment[s][a] / st * 100) if st else Decimal(0) for a in (2, 3))
    gerekce_say = defaultdict(lambda: [0, Decimal(0)])
    for k in krediler:
        for g in k["gerekce"]:
            ad = g.split(" (")[0].split(" kademe")[0]
            ad = "Gecikme" if ad.startswith("Gecikme") else "Not kötüleşmesi" if ad.startswith("Not") else "TO artışı" if ad.startswith("TO") else "Müşteri bazı" if ad.startswith("Müşteri") else ad
            gerekce_say[(k["asama"], ad)][0] += 1
            gerekce_say[(k["asama"], ad)][1] += k["bakiye"]
    r = {"krediler": krediler, "ozet": ozet, "toplam": toplam, "segment": segment, "gerekce": dict(gerekce_say), "uyarilar": uy, "gecis": None, "degisen": []}
    eski = {k["no"]: k["onceki_asama"] for k in krediler if k["onceki_asama"]}
    eski_bakiye = {}
    if onceki is not None:
        for o in onceki:
            o["asama"] = o["onceki_asama"] or siniflandir(o, **esik)[0]              # önceki dosyada Aşama sütunu varsa o kullanılır
        eski, eski_bakiye = {o["no"]: o["asama"] for o in onceki}, {o["no"]: o["bakiye"] for o in onceki}
    if eski:
        m = {a: {b: Decimal(0) for b in (1, 2, 3, "Kapandı")} for a in (1, 2, 3, "Yeni")}
        yeni_no = {k["no"] for k in krediler}
        for k in krediler:
            e = k["onceki_asama"] = eski.get(k["no"])
            m[e or "Yeni"][k["asama"]] += eski_bakiye.get(k["no"], k["bakiye"]) if e else k["bakiye"]
            if e and e != k["asama"]:
                r["degisen"].append({"k": k, "onceki": e, "yon": "Kötüleşti" if k["asama"] > e else "İyileşti"})
                if k["asama"] < e:
                    uy.append(uyari("Orta", f"Aşama {e} → {k['asama']} iyileşme", k["no"], "Üst aşamadan çıkış için politikanızdaki izleme süresi ve ödeme performansı "
                                                                                         "koşullarının sağlandığını doğrulayın."))
        for no, a in eski.items():
            if no not in yeni_no:
                m[a]["Kapandı"] += eski_bakiye.get(no, Decimal(0))
        r["gecis"] = m
        r["degisen"].sort(key=lambda d: (d["yon"] != "Kötüleşti", -d["k"]["asama"], -d["k"]["bakiye"]))
    return r


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

ASAMA_RENK = {1: YESIL, 2: SARI, 3: KIRMIZI}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    oz = wb.active
    oz.title = "Özet"
    _baslik(oz, ["Aşama", "Kredi Adedi", "Risk Bakiyesi", "Pay %"], (26, 12, 20, 9))
    for a in (1, 2, 3):
        _satir(oz, [f"Aşama {a}", r["ozet"][a]["adet"], r["ozet"][a]["bakiye"], r["ozet"][a]["pay"]], para_s=(3, 4), dolgu={1: ASAMA_RENK[a]})
    _satir(oz, ["Toplam", len(r["krediler"]), r["toplam"], 100], para_s=(3,))
    oz.append([])
    _baslik(oz, ["Segment", "Kredi Adedi", "Risk Bakiyesi", "Aşama 1", "Aşama 2", "Aşama 3", "Aşama 2 Payı %", "Aşama 3 Payı %"], (26, 12, 20, 20, 20, 20, 12, 12), None)
    for s, v in r["segment"].items():
        _satir(oz, [s, v["adet"], v["bakiye"], v[1], v[2], v[3], v["a2_pay"], v["a3_pay"]], para_s=(3, 4, 5, 6))
    oz.append([])
    _baslik(oz, ["Gerekçe", "Aşama", "Kredi Adedi", "Risk Bakiyesi"], dondur=None)
    for (a, g), (adet, bak) in sorted(r["gerekce"].items(), key=lambda x: (-x[0][0], -x[1][1])):
        _satir(oz, [g, a, adet, bak], para_s=(4,))
    oz.append([])
    oz.append(["Bir kredi birden çok gerekçeyle sayılabilir. Bu tablo sınıflandırma önerisidir; beklenen kredi zararı (karşılık) hesaplanmaz."])

    sn = wb.create_sheet("Sınıflandırma")
    _baslik(sn, ["Kredi No", "Müşteri No", "Müşteri", "Segment", "Ürün", "Risk Bakiyesi", "Gecikme Gün", "Aşama", "Önceki Aşama", "Gerekçe", "Değerlendirme Notu"],
            (14, 12, 28, 12, 18, 18, 9, 8, 9, 70, 30))
    for k in sorted(r["krediler"], key=lambda k: (-k["asama"], -k["bakiye"])):
        _satir(sn, [k["no"], k["musteri_no"], k["musteri"], k["segment"], k["urun"], k["bakiye"], k["gun"], k["asama"], k["onceki_asama"], "; ".join(k["gerekce"]), ""], para_s=(6,),
               dolgu={8: ASAMA_RENK[k["asama"]], 11: SARI}, sar=(10,))
    sn.auto_filter.ref = f"A1:K{sn.max_row}"
    if r["gecis"] is not None:
        dg = wb.create_sheet("Aşama Değişenler")
        _baslik(dg, ["Yön", "Kredi No", "Müşteri", "Segment", "Risk Bakiyesi", "Önceki Aşama", "Yeni Aşama", "Gerekçe"], (12, 14, 28, 12, 18, 9, 9, 70))
        for d in r["degisen"]:
            k = d["k"]
            _satir(dg, [d["yon"], k["no"], k["musteri"], k["segment"], k["bakiye"], d["onceki"], k["asama"], "; ".join(k["gerekce"]) or "Aşama 2 / 3 ölçütü kalmadı"], para_s=(5,),
                   dolgu={1: KIRMIZI if d["yon"] == "Kötüleşti" else YESIL}, sar=(8,))
        gm = wb.create_sheet("Geçiş Matrisi")
        _baslik(gm, ["Önceki ↓ / Bu Dönem →", "Aşama 1", "Aşama 2", "Aşama 3", "Kapandı", "Toplam"], (24, 20, 20, 20, 20, 20))
        for a in (1, 2, 3, "Yeni"):
            s = r["gecis"][a]
            _satir(gm, [f"Aşama {a}" if a != "Yeni" else "Yeni kredi", s[1], s[2], s[3], s["Kapandı"], sum(s.values())], para_s=(2, 3, 4, 5, 6))
        gm.append([])
        gm.append(["Tutarlar önceki dönem bakiyesidir (önceki dosya verilmediyse veya yeni kredide bu dönem bakiyesi)."])
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, onceki_yol: Path | None = None, olcek: list[str] | None = None, bulasma: set[str] = frozenset(), **esik) -> dict:
    olcek = olcek or []
    krediler, uy = oku(yol, olcek)
    if not krediler:
        raise ValueError(f"{yol.name}: okunabilir kredi yok")
    onceki = None
    if onceki_yol:
        onceki, uy2 = oku(onceki_yol, olcek)
        uy += [u for u in uy2 if u["onem"] == "Yüksek"]
    r = analiz(krediler, onceki, bulasma, **esik)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Kredileri gecikme günü ve risk göstergelerine göre TFRS 9 Aşama 1-2-3 olarak sınıflandırır.")
    p.add_argument("--krediler", type=Path, default=ORNEK / "krediler.csv",
                   help="Kredi No, Müşteri No, Müşteri, Segment, Ürün, Risk Bakiyesi, Gecikme Gün, Yeniden Yapılandırma, Yakın İzleme, Temerrüt, İlk Not, Güncel Not, İlk TO, Güncel TO")
    p.add_argument("--onceki", type=Path, help="Önceki dönem dosyası (aynı biçim; 'Aşama' sütunu varsa o kullanılır)")
    p.add_argument("--asama2-gun", type=int, default=30, help="Aşama 2 gecikme eşiği (varsayılan 30)")
    p.add_argument("--temerrut-gun", type=int, default=90, help="Aşama 3 gecikme eşiği (varsayılan 90)")
    p.add_argument("--not-esik", type=int, default=3, help="Aşama 2 için not kötüleşmesi, kademe (varsayılan 3)")
    p.add_argument("--to-kat", type=Decimal, default=Decimal(2), help="Aşama 2 için güncel TO ÷ ilk TO (varsayılan 2)")
    p.add_argument("--not-olcegi", default="", help='Notlar en iyiden en kötüye, virgülle (ör. "AAA,AA,A,BBB,BB,B,CCC"); boşsa notlar sayısaldır')
    p.add_argument("--bulasma", default="", help='Müşteri bazında Aşama 3 uygulanacak segmentler (ör. "Ticari,KOBİ")')
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "tfrs9_asama.xlsx")
    a = p.parse_args(argv)
    ornek = a.krediler == ORNEK / "krediler.csv"
    onceki = a.onceki or (ORNEK / "onceki_donem.csv" if ornek else None)
    for y in (a.krediler, onceki):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    if not (0 < a.asama2_gun < a.temerrut_gun) or a.not_esik < 1 or a.to_kat <= 1:
        print("[X] 0 < --asama2-gun < --temerrut-gun; --not-esik en az 1; --to-kat 1'den büyük olmalı")
        return 2
    olcek = [x.strip().upper() for x in (a.not_olcegi or ("AAA,AA,A,BBB,BB,B,CCC,CC,C" if ornek else "")).split(",") if x.strip()]
    bulasma = {katla(x) for x in (a.bulasma or ("Ticari,KOBİ" if ornek else "")).split(",") if x.strip()}
    try:
        r = calistir(a.krediler, a.cikti, onceki, olcek, bulasma, a2_gun=a.asama2_gun, a3_gun=a.temerrut_gun, not_esik=a.not_esik, to_kat=a.to_kat)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(r['krediler'])} kredi · " + " · ".join(f"Aşama {x}: {r['ozet'][x]['adet']} (%{r['ozet'][x]['pay']})" for x in (1, 2, 3)))
    if r["gecis"] is not None:
        print(f"[OK] Önceki döneme göre: {sum(d['yon'] == 'Kötüleşti' for d in r['degisen'])} kötüleşen, {sum(d['yon'] == 'İyileşti' for d in r['degisen'])} iyileşen")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['kim']} · {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
