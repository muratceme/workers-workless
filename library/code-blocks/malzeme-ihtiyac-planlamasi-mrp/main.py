"""
Malzeme İhtiyaç Planlaması (MRP) — Workers / Workless kod bloğu
Üretim › Üretim Planlama › Üretim Planlama Uzmanı

Ana üretim planından (MPS) ve çok seviyeli ürün ağacından (BOM) her malzeme için dönem bazında standart MRP
kaydını hesaplar:

    Brüt ihtiyaç → (eldeki + planlanmış girişler − emniyet stoğu) → net ihtiyaç → planlanan sipariş girişi
    (lot kuralına göre) → tedarik süresi kadar öne alınmış planlanan sipariş verilişi

Bir üst malzemenin planlanan sipariş verilişi, alt malzemenin aynı dönemdeki brüt ihtiyacını oluşturur
(adet × (1 + fire %)). Malzemeler düşük seviye koduna göre işlenir; bir malzeme ağacın birden fazla
seviyesinde geçse de bir kez, en alt seviyesinde hesaplanır. Lot kuralları: sipariş bazında (L4L), sabit lot,
en az sipariş miktarı ve ambalaj katı. Verilişi planlama ufkunun başından önceye düşen siparişler "GEÇMİŞ"
olarak işaretlenir (hemen verilmeli / acil). İnternete bağlanmaz.

Kullanım:
    python main.py                                       # örnek veriyle dener
    python main.py --mps mps.xlsx --bom urun_agaci.xlsx --stok stok.xlsx --baslangic 02.11.2026 --donem 8
"""
from __future__ import annotations

import argparse
import csv
import math
import sys
from collections import OrderedDict, defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def sayi(x, varsayilan: float = 0.0) -> float:
    if x in (None, ""):
        return varsayilan
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip().replace("%", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(Decimal(s))
    except InvalidOperation:
        return varsayilan


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(str(x).strip(), f).date()
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def sutunlar(baslik, alanlar):
    b = [kucuk(x) for x in baslik]
    return {a: next((i for i, x in enumerate(b) if x in adlar), None) for a, adlar in alanlar.items()}


class Donemler:
    """Dönem 1 = başlangıç gününden itibaren 'gun' günlük kova (varsayılan haftalık)."""

    def __init__(self, baslangic: date, sayi_: int, gun: int = 7):
        self.baslangic, self.sayi, self.gun = baslangic, sayi_, gun

    def bul(self, x) -> int | None:
        if isinstance(x, (int, float)) or (isinstance(x, str) and x.strip().isdigit()):
            return int(float(x))
        t = tarih(x)
        if t is None:
            return None
        return (t - self.baslangic).days // self.gun + 1

    def etiket(self, d: int) -> str:
        return f"{d}. dönem ({self.baslangic + timedelta(days=(d - 1) * self.gun):%d.%m})"


# ----------------------------------------------------------------------------
# Okuma
# ----------------------------------------------------------------------------

def oku(mps_yolu: Path, bom_yolu: Path, stok_yolu: Path, donem: Donemler):
    uyarilar = []
    s = tablo_oku(mps_yolu)
    k = sutunlar(s[0], {"kod": ("ürün", "ürün kodu", "malzeme", "stok kodu"), "donem": ("dönem", "hafta", "tarih", "termin"),
                        "miktar": ("miktar", "adet")})
    if None in k.values():
        raise SystemExit(f"MPS'de Ürün, Dönem (veya Tarih) ve Miktar gerekli. Başlıklar: {s[0]}")
    mps: dict[str, dict[int, float]] = defaultdict(lambda: defaultdict(float))
    for r in s[1:]:
        d = donem.bul(r[k["donem"]])
        if d is None:
            continue
        if d < 1:
            uyarilar.append(f"MPS {r[k['kod']]}: {r[k['donem']]} planlama ufkunun başından önce; 1. döneme alındı")
            d = 1
        if d > donem.sayi:
            uyarilar.append(f"MPS {r[k['kod']]}: {r[k['donem']]} ufkun dışında ({donem.sayi} dönem); dikkate alınmadı")
            continue
        mps[str(r[k["kod"]]).strip()][d] += sayi(r[k["miktar"]])

    s = tablo_oku(bom_yolu)
    k = sutunlar(s[0], {"ust": ("üst malzeme", "ana malzeme", "mamul", "üst"), "alt": ("alt malzeme", "bileşen", "alt"),
                        "miktar": ("miktar", "birim miktar", "kullanım miktarı"), "fire": ("fire %", "fire", "fire oranı")})
    if None in (k["ust"], k["alt"], k["miktar"]):
        raise SystemExit(f"Ürün ağacında Üst Malzeme, Alt Malzeme ve Miktar gerekli. Başlıklar: {s[0]}")
    bom: dict[str, list[tuple[str, float, float]]] = defaultdict(list)
    for r in s[1:]:
        if r[k["ust"]] in (None, ""):
            continue
        fire = sayi(r[k["fire"]]) if k["fire"] is not None else 0.0
        bom[str(r[k["ust"]]).strip()].append((str(r[k["alt"]]).strip(), sayi(r[k["miktar"]]), fire / 100 if fire >= 1 else fire))

    s = tablo_oku(stok_yolu)
    k = sutunlar(s[0], {"kod": ("malzeme", "stok kodu", "ürün", "malzeme kodu"), "ad": ("malzeme adı", "ad", "açıklama"),
                        "eldeki": ("eldeki", "eldeki stok", "stok"), "emniyet": ("emniyet stoğu", "emniyet"),
                        "lt": ("tedarik süresi (dönem)", "tedarik süresi", "temin süresi"), "lot": ("lot yöntemi", "lot kuralı"),
                        "lot_miktar": ("lot miktarı", "sabit lot"), "moq": ("en az sipariş", "moq"), "kat": ("ambalaj katı", "sipariş katı"),
                        "tur": ("tedarik türü", "tür", "üret/satın al"), "birim": ("birim",)})
    if k["kod"] is None or k["lt"] is None:
        raise SystemExit(f"Stok/malzeme listesinde Malzeme ve Tedarik Süresi gerekli. Başlıklar: {s[0]}")
    al = lambda r, a: r[k[a]] if k[a] is not None and k[a] < len(r) else None  # noqa: E731
    kalemler = OrderedDict()
    for r in s[1:]:
        kod = str(al(r, "kod") or "").strip()
        if kod:
            kalemler[kod] = {"ad": str(al(r, "ad") or ""), "eldeki": sayi(al(r, "eldeki")), "emniyet": sayi(al(r, "emniyet")),
                             "lt": int(sayi(al(r, "lt"))), "lot": kucuk(al(r, "lot")) or "l4l", "lot_miktar": sayi(al(r, "lot_miktar")),
                             "moq": sayi(al(r, "moq")), "kat": sayi(al(r, "kat")), "tur": str(al(r, "tur") or ""), "birim": str(al(r, "birim") or ""),
                             "acik": defaultdict(float)}
    return mps, bom, kalemler, uyarilar


def acik_siparisleri_oku(yol: Path | None, kalemler: dict, donem: Donemler, uyarilar: list[str]) -> None:
    if not yol:
        return
    s = tablo_oku(yol)
    k = sutunlar(s[0], {"kod": ("malzeme", "stok kodu", "ürün"), "miktar": ("miktar", "açık miktar", "kalan"),
                        "donem": ("dönem", "termin", "teslim tarihi", "tarih")})
    if None in k.values():
        raise SystemExit(f"Açık siparişlerde Malzeme, Miktar ve Termin gerekli. Başlıklar: {s[0]}")
    for r in s[1:]:
        kod = str(r[k["kod"]] or "").strip()
        d = donem.bul(r[k["donem"]])
        if kod in kalemler and d is not None:
            if d < 1:
                uyarilar.append(f"{kod}: açık sipariş termini geçmiş ({r[k['donem']]}); 1. dönemde gelecek kabul edildi — tedarikçiyle teyit edin")
                d = 1
            if d <= donem.sayi:
                kalemler[kod]["acik"][d] += sayi(r[k["miktar"]])


# ----------------------------------------------------------------------------
# Hesap
# ----------------------------------------------------------------------------

def dusuk_seviye_kodlari(bom: dict, kokler: set[str]) -> dict[str, int]:
    seviye: dict[str, int] = {}

    def gez(kod: str, s: int, yol: tuple):
        if kod in yol:
            raise SystemExit(f"Ürün ağacında döngü: {' → '.join(yol + (kod,))}")
        if seviye.get(kod, -1) < s:
            seviye[kod] = s
        for alt, _, _ in bom.get(kod, []):
            gez(alt, s + 1, yol + (kod,))

    for kok in kokler:
        gez(kok, 0, ())
    return seviye


def lot(miktar: float, k: dict) -> float:
    if miktar <= 0:
        return 0.0
    q = miktar
    if k["lot"] in ("sabit", "sabit lot", "foq") and k["lot_miktar"] > 0:
        q = math.ceil(q / k["lot_miktar"] - 1e-9) * k["lot_miktar"]
    q = max(q, k["moq"])
    if k["kat"] > 0:
        q = math.ceil(q / k["kat"] - 1e-9) * k["kat"]
    return q


def mrp(mps, bom, kalemler, donem: Donemler) -> dict:
    uyarilar = []
    kokler = set(mps)
    for kod in list(kokler) + [a for v in bom.values() for a, _, _ in v]:
        if kod not in kalemler:
            uyarilar.append(f"{kod}: malzeme listesinde yok; stok 0, tedarik süresi 0, L4L kabul edildi")
            kalemler[kod] = {"ad": "", "eldeki": 0.0, "emniyet": 0.0, "lt": 0, "lot": "l4l", "lot_miktar": 0, "moq": 0, "kat": 0, "tur": "",
                             "birim": "", "acik": defaultdict(float)}
    llc = dusuk_seviye_kodlari(bom, kokler)
    brut: dict[str, dict[int, float]] = defaultdict(lambda: defaultdict(float))
    for kod, v in mps.items():
        for d, q in v.items():
            brut[kod][d] += q
    kayit = OrderedDict()
    gecmis = []
    for kod in sorted(llc, key=lambda x: (llc[x], x)):
        k = kalemler[kod]
        n = donem.sayi
        satir = {"brut": [0.0] * (n + 1), "giris": [0.0] * (n + 1), "eldeki": [0.0] * (n + 1), "net": [0.0] * (n + 1),
                 "pgiris": [0.0] * (n + 1), "pveris": [0.0] * (n + 1)}
        eldeki = k["eldeki"]
        if eldeki < k["emniyet"]:
            uyarilar.append(f"{kod}: başlangıç stoğu ({eldeki:g}) emniyet stoğunun ({k['emniyet']:g}) altında")
        for d in range(1, n + 1):
            satir["brut"][d] = brut[kod].get(d, 0.0)
            satir["giris"][d] = k["acik"].get(d, 0.0)
            mevcut = eldeki + satir["giris"][d] - satir["brut"][d]
            if mevcut < k["emniyet"]:
                net = k["emniyet"] - mevcut
                satir["net"][d] = net
                q = lot(net, k)
                satir["pgiris"][d] = q
                mevcut += q
                vd = d - k["lt"]
                if vd < 1:
                    gecmis.append((kod, d, q, vd))
                    satir["pveris"][1] += q                       # hemen verilmeli
                else:
                    satir["pveris"][vd] += q
            satir["eldeki"][d] = mevcut
            eldeki = mevcut
        # Alt malzemelerin brüt ihtiyacı: planlanan veriliş × miktar × (1 + fire)
        for alt, miktar, fire in bom.get(kod, []):
            for d in range(1, n + 1):
                if satir["pveris"][d]:
                    brut[alt][d] += satir["pveris"][d] * miktar * (1 + fire)
        kayit[kod] = {"satir": satir, "llc": llc[kod], "kalem": k}
    return {"kayit": kayit, "gecmis": gecmis, "uyarilar": uyarilar}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
GRI = PatternFill("solid", fgColor="EDEDED")
SARI = PatternFill("solid", fgColor="FFF4CE")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SAYI = "#,##0.##"
SATIRLAR = [("brut", "Brüt ihtiyaç"), ("giris", "Planlanmış girişler (açık sipariş)"), ("eldeki", "Öngörülen eldeki stok"),
            ("net", "Net ihtiyaç"), ("pgiris", "Planlanan sipariş girişi"), ("pveris", "Planlanan sipariş verilişi")]


def calistir(mps_yolu: Path, bom_yolu: Path, stok_yolu: Path, cikti: Path, baslangic: date, donem_sayisi: int = 8, donem_gun: int = 7,
             acik_yolu: Path | None = None) -> dict:
    donem = Donemler(baslangic, donem_sayisi, donem_gun)
    mps, bom, kalemler, uyarilar = oku(mps_yolu, bom_yolu, stok_yolu, donem)
    acik_siparisleri_oku(acik_yolu, kalemler, donem, uyarilar)
    sonuc = mrp(mps, bom, kalemler, donem)
    sonuc["uyarilar"] = uyarilar + sonuc["uyarilar"]
    n = donem_sayisi

    wb = Workbook()
    ws = wb.active
    ws.title = "MRP Kayıtları"
    ws.append(["Malzeme", "Seviye", "Satır"] + [donem.etiket(d) for d in range(1, n + 1)])
    for c in ws[1]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI
    for kod, x in sonuc["kayit"].items():
        k = x["kalem"]
        ws.append([f"{kod} {k['ad']}".strip(), x["llc"], f"Eldeki {k['eldeki']:g} · ES {k['emniyet']:g} · TS {k['lt']} dönem · lot {k['lot'].upper()}"
                   + (f" {k['lot_miktar']:g}" if k["lot_miktar"] else "") + (f" · MOQ {k['moq']:g}" if k["moq"] else "") + (f" · kat {k['kat']:g}" if k["kat"] else "")])
        for c in ws[ws.max_row]:
            c.fill = GRI
            c.font = Font(bold=True)
        for alan, ad in SATIRLAR:
            ws.append(["", "", ad] + [x["satir"][alan][d] or None for d in range(1, n + 1)])
            for j in range(4, 4 + n):
                ws.cell(ws.max_row, j).number_format = SAYI
                v = ws.cell(ws.max_row, j).value
                if alan == "pveris" and v:
                    ws.cell(ws.max_row, j).fill = SARI
                if alan == "eldeki" and v is not None and v < k["emniyet"]:
                    ws.cell(ws.max_row, j).fill = KIRMIZI
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 7
    ws.column_dimensions["C"].width = 40
    for j in range(4, 4 + n):
        ws.column_dimensions[get_column_letter(j)].width = 15
    ws.freeze_panes = "D2"

    s = wb.create_sheet("Sipariş Önerileri", 0)
    s.append(["Malzeme", "Ad", "Tedarik Türü", "Miktar", "Birim", "Sipariş Verilişi", "İhtiyaç (Giriş) Dönemi", "Durum"])
    for c in s[1]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI
    gecmis = {(g[0], g[1]) for g in sonuc["gecmis"]}
    for kod, x in sonuc["kayit"].items():
        k = x["kalem"]
        for d in range(1, n + 1):
            q = x["satir"]["pgiris"][d]
            if q:
                vd = max(1, d - k["lt"])
                acil = (kod, d) in gecmis
                s.append([kod, k["ad"], k["tur"], q, k["birim"], "HEMEN (GEÇMİŞ)" if acil else donem.etiket(vd), donem.etiket(d),
                          "Acil: tedarik süresi yetmiyor" if acil else ("Bu dönem verilmeli" if vd == 1 else "Planlı")])
                s.cell(s.max_row, 4).number_format = SAYI
                if acil:
                    s.cell(s.max_row, 8).fill = KIRMIZI
                elif vd == 1:
                    s.cell(s.max_row, 8).fill = SARI
    for j, w in enumerate((16, 28, 13, 12, 8, 20, 20, 30), 1):
        s.column_dimensions[get_column_letter(j)].width = w
    s.auto_filter.ref = s.dimensions
    for u in sonuc["uyarilar"]:
        s.append(["Uyarı", u])

    b = wb.create_sheet("Bilgi")
    for satir in [["Dönem", f"1. dönem {baslangic:%d.%m.%Y}'den başlar, her dönem {donem_gun} gün"],
                  ["Net ihtiyaç", "Öngörülen stok (önceki eldeki + planlanmış giriş − brüt ihtiyaç) emniyet stoğunun altına düşerse aradaki fark"],
                  ["Lot kuralı", "L4L: net ihtiyaç kadar · SABİT: lot miktarının katı · ardından en az sipariş ve ambalaj katına yukarı yuvarlama"],
                  ["Alt malzeme", "Üst malzemenin planlanan verilişi × birim miktar × (1 + fire oranı) = alt malzemenin aynı dönemdeki brüt ihtiyacı"],
                  ["Düşük seviye kodu", "Bir malzeme ağaçta birden fazla seviyede geçerse en alt seviyesinde, tüm ihtiyaçlar toplanınca bir kez hesaplanır"],
                  ["GEÇMİŞ", "Veriliş dönemi ufkun başından önceye düşen sipariş: hemen verilmeli, tedarik süresini kısaltma veya MPS'i kaydırma gerekir"],
                  ["Not", "Kapasite kısıtı yoktur (sonsuz kapasite MRP); üretim kapasitesi için Haftalık Üretim Çizelgesi kullanın"]]:
        b.append(satir)
    b.column_dimensions["A"].width = 18
    b.column_dimensions["B"].width = 130
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    return sonuc


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description="MPS ve ürün ağacından dönem bazında MRP kayıtları ve sipariş önerileri üretir.")
    ap.add_argument("--mps", type=Path, default=BURASI / "ornek_veri" / "mps.csv", help="Ana üretim planı: Ürün, Dönem (no veya tarih), Miktar")
    ap.add_argument("--bom", type=Path, default=BURASI / "ornek_veri" / "urun_agaci.csv", help="Ürün ağacı: Üst Malzeme, Alt Malzeme, Miktar, Fire %%")
    ap.add_argument("--stok", type=Path, default=BURASI / "ornek_veri" / "malzemeler.csv",
                    help="Malzemeler: Malzeme, Eldeki, Emniyet Stoğu, Tedarik Süresi (dönem), Lot Yöntemi, Lot Miktarı, En Az Sipariş, Ambalaj Katı")
    ap.add_argument("--acik", type=Path, help="Açık siparişler (planlanmış girişler): Malzeme, Miktar, Termin")
    ap.add_argument("--baslangic", help="1. dönemin başlangıcı (GG.AA.YYYY)")
    ap.add_argument("--donem", type=int, default=8, help="Dönem sayısı (varsayılan 8)")
    ap.add_argument("--donem-gun", type=int, default=7, help="Dönem uzunluğu, gün (varsayılan 7 = haftalık)")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "mrp.xlsx")
    a = ap.parse_args(argv)
    if a.mps == BURASI / "ornek_veri" / "mps.csv":
        a.acik = a.acik or BURASI / "ornek_veri" / "acik_siparisler.csv"
        a.baslangic = a.baslangic or "02.11.2026"
    bas = tarih(a.baslangic) if a.baslangic else date.today() - timedelta(days=date.today().weekday())
    s = calistir(a.mps, a.bom, a.stok, a.cikti, bas, a.donem, a.donem_gun, a.acik)
    oneriler = sum(1 for x in s["kayit"].values() for q in x["satir"]["pgiris"] if q)
    print(f"[OK] {len(s['kayit'])} malzeme · {oneriler} sipariş önerisi · geçmiş (acil): {len(s['gecmis'])}")
    for kod, d, q, vd in s["gecmis"]:
        print(f"[!] {kod}: {d}. dönem ihtiyacı için {q:g} birim siparişin {1 - vd} dönem önce verilmesi gerekirdi")
    for u in s["uyarilar"]:
        print(f"[!] {u}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
