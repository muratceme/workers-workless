"""
Performans Primi Hesaplama — Workers / Workless kod bloğu
İnsan Kaynakları › Ücretlendirme ve Yan Haklar Uzmanı

Hedef gerçekleşmelerinden ve prim skalasından çalışan bazında brüt performans primini hesaplar:
  - Hedef gerçekleşme: Artan hedefte gerçekleşen ÷ hedef, Azalan hedefte (fire, süre, şikâyet) hedef ÷ gerçekleşen.
    Hedef 0 olan Azalan hedefte (ör. iş kazası) gerçekleşen 0 ise %100, değilse %0. Her hedef --hedef-tavan ile
    sınırlanır (varsayılan %150).
  - Ağırlıklı skor = Σ ağırlık × gerçekleşme. Ağırlık toplamı 100 değilse orantılanır ve uyarı verilir.
  - Prim çarpanı skaladan okunur: ilk satırın altı 0 (eşik), satırlar arası doğrusal (veya --kademeli),
    son satırın üstü son çarpan (tavan).
  - İsteğe bağlı şirket çarpanı: çarpan = bireysel ağırlık × bireysel çarpan + (1 − bireysel ağırlık) × şirket çarpanı.
  - Hedef prim = brüt ücret × maaş katı (veya doğrudan TL). Dönem içinde giren / ayrılan çalışanda kıst
    (çalışılan gün ÷ dönem günü). Dönem bitmeden ayrılana varsayılan olarak ödenmez (--ayrilana-ode).
Rapor: prim listesi, hedef detayı, departman özeti, çarpan dağılımı, uyarılar. İnternete bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: 8 çalışan, 2026 1. yarıyıl
    python main.py --calisanlar c.xlsx --hedefler h.xlsx --skala skala.csv --donem 01.01.2026 30.06.2026
    python main.py --calisanlar c.xlsx --hedefler h.xlsx --sirket-carpani 110 --bireysel-agirlik 70 --butce 900000
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
K2 = Decimal("0.01")
YUZ = Decimal(100)

CALISAN_SUTUNLARI = {"sicil": ("sicil no", "sicil"), "ad": ("ad soyad", "ad"), "departman": ("departman", "birim"), "pozisyon": ("pozisyon", "unvan"),
                     "ucret": ("brut ucret tl", "brut ucret", "ucret", "aylik brut"),
                     "kat": ("hedef prim maas", "hedef prim maas kati", "prim maas kati", "maas kati"),
                     "hedef_tl": ("hedef prim tl", "hedef prim tutari", "hedef prim"),
                     "giris": ("ise giris", "ise giris tarihi"), "cikis": ("ayrilis tarihi", "cikis tarihi", "isten cikis")}
HEDEF_SUTUNLARI = {"sicil": ("sicil no", "sicil"), "hedef": ("hedef", "hedef adi", "kpi", "gosterge"), "agirlik": ("agirlik", "agirlik"),
                   "deger": ("hedef deger", "hedef degeri", "hedeflenen"), "gercek": ("gerceklesen", "gerceklesme degeri", "fiili"),
                   "oran": ("gerceklesme", "gerceklesme orani", "basari orani"), "yon": ("yon", "yonu", "hedef yonu")}
SKALA_SUTUNLARI = {"g": ("gerceklesme", "gerceklesme orani", "skor", "puan"), "c": ("prim carpani", "carpan", "odeme orani")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def sayi(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("%", "").replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tl(x: Decimal | None) -> str:
    return "—" if x is None else f"{x.quantize(K2, ROUND_HALF_UP):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def yuzde(x: Decimal | None) -> str:
    return "—" if x is None else f"%{x * 100:.1f}".replace(".", ",")


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Hedef:
    ad: str
    agirlik: Decimal
    deger: Decimal | None
    gercek: Decimal | None
    yon: str                   # "artan" | "azalan"
    oran: Decimal | None = None          # gerçekleşme (1 = %100), tavan öncesi
    sinirli: Decimal | None = None       # tavan sonrası
    satir: int = 0


@dataclass
class Calisan:
    sicil: str
    ad: str
    departman: str
    pozisyon: str
    ucret: Decimal | None
    kat: Decimal | None
    hedef_tl: Decimal | None
    giris: date | None
    cikis: date | None
    hedefler: list = field(default_factory=list)
    skor: Decimal | None = None
    bireysel: Decimal | None = None
    carpan: Decimal | None = None
    kist: Decimal = Decimal(1)
    gun: int = 0
    prim: Decimal | None = None
    notlar: list = field(default_factory=list)

    @property
    def hedef_prim(self) -> Decimal | None:
        if self.hedef_tl is not None:
            return self.hedef_tl
        if self.ucret is not None and self.kat is not None:
            return (self.ucret * self.kat).quantize(K2, ROUND_HALF_UP)
        return None


def oku_calisanlar(yol: Path) -> list[Calisan]:
    sonuc = []
    for r in kayitlar(yol, CALISAN_SUTUNLARI, ("sicil",)):
        if not r.get("sicil"):
            continue
        sonuc.append(Calisan(metin(r["sicil"]), metin(r.get("ad")), metin(r.get("departman")) or "—", metin(r.get("pozisyon")), sayi(r.get("ucret")),
                             sayi(r.get("kat")), sayi(r.get("hedef_tl")), tarih(r.get("giris")), tarih(r.get("cikis"))))
    return sonuc


def oku_hedefler(yol: Path) -> dict[str, list[Hedef]]:
    g: dict[str, list[Hedef]] = defaultdict(list)
    for r in kayitlar(yol, HEDEF_SUTUNLARI, ("sicil", "agirlik")):
        if not r.get("sicil"):
            continue
        yon = "azalan" if katla(r.get("yon")).startswith(("azal", "dusuk", "min")) else "artan"
        h = Hedef(metin(r.get("hedef")) or f"Hedef (satır {r['_satir']})", sayi(r.get("agirlik")) or Decimal(0), sayi(r.get("deger")),
                  sayi(r.get("gercek")), yon, satir=r["_satir"])
        dogrudan = sayi(r.get("oran"))
        if dogrudan is not None and (h.deger is None or h.gercek is None):
            h.oran = dogrudan / YUZ if dogrudan > 3 else dogrudan
        g[metin(r["sicil"])].append(h)
    return g


def oku_skala(yol: Path) -> list[tuple[Decimal, Decimal]]:
    noktalar = []
    for r in kayitlar(yol, SKALA_SUTUNLARI, ("g", "c")):
        g, c = sayi(r.get("g")), sayi(r.get("c"))
        if g is not None and c is not None:
            noktalar.append((g / YUZ, c / YUZ))
    if not noktalar:
        raise ValueError(f"{yol.name}: skala satırı yok")
    return sorted(noktalar)


# ----------------------------------------------------------------------------
# Hesap
# ----------------------------------------------------------------------------

def gerceklesme(h: Hedef) -> Decimal | None:
    if h.oran is not None:
        return h.oran
    if h.deger is None or h.gercek is None:
        return None
    if h.yon == "azalan":
        if h.deger == 0:
            return Decimal(1) if h.gercek <= 0 else Decimal(0)
        if h.gercek <= 0:
            return None                          # tavan uygulanacak; aşağıda ele alınır
        return h.deger / h.gercek
    if h.deger == 0:
        return None
    return max(h.gercek / h.deger, Decimal(0))


def carpan_bul(skor: Decimal, skala: list[tuple[Decimal, Decimal]], kademeli: bool = False) -> Decimal:
    if skor < skala[0][0]:
        return Decimal(0)
    if skor >= skala[-1][0]:
        return skala[-1][1]
    for (g1, c1), (g2, c2) in zip(skala, skala[1:]):
        if g1 <= skor < g2:
            return c1 if kademeli else c1 + (c2 - c1) * (skor - g1) / (g2 - g1)
    return skala[-1][1]


def hesapla(calisanlar: list[Calisan], hedefler: dict[str, list[Hedef]], skala, *, donem: tuple[date, date], hedef_tavan: Decimal = Decimal("1.5"),
            kademeli: bool = False, sirket: Decimal | None = None, bireysel_agirlik: Decimal = Decimal(1), ayrilana_ode: bool = False,
            min_gun: int = 0) -> list[dict]:
    uyarilar = []

    def uyar(onem, tur, c, aciklama):
        uyarilar.append({"onem": onem, "tur": tur, "sicil": c.sicil if c else "", "aciklama": aciklama})
        if c:
            c.notlar.append(tur)

    bas, bit = donem
    donem_gun = (bit - bas).days + 1
    sicil_set = {c.sicil for c in calisanlar}
    for s in sorted(set(hedefler) - sicil_set):
        uyarilar.append({"onem": "Orta", "tur": "Çalışan listesinde yok", "sicil": s, "aciklama": f"{len(hedefler[s])} hedef satırı var ama çalışan listesinde yok"})
    for c in calisanlar:
        c.hedefler = hedefler.get(c.sicil, [])
        # Kıst
        g1 = max(bas, c.giris) if c.giris else bas
        g2 = min(bit, c.cikis) if c.cikis else bit
        c.gun = max((g2 - g1).days + 1, 0)
        c.kist = Decimal(c.gun) / Decimal(donem_gun)
        if c.giris and c.giris > bit:
            c.kist, c.gun = Decimal(0), 0
        if not c.hedefler:
            uyar("Yüksek", "Hedef yok", c, "Hedef kaydı yok; prim hesaplanmadı")
            continue
        toplam_agirlik = sum((h.agirlik for h in c.hedefler), Decimal(0))
        if toplam_agirlik <= 0:
            uyar("Yüksek", "Ağırlık hatası", c, "Hedef ağırlıkları toplamı 0")
            continue
        if toplam_agirlik != YUZ:
            uyar("Orta", "Ağırlık toplamı", c, f"Ağırlık toplamı {toplam_agirlik:g}; 100'e orantılandı")
        skor, eksik = Decimal(0), False
        for h in c.hedefler:
            g = gerceklesme(h)
            if g is None:
                if h.yon == "azalan" and h.gercek is not None and h.gercek <= 0 and h.deger:
                    g = hedef_tavan
                else:
                    uyar("Yüksek", "Hedef verisi eksik", c, f"'{h.ad}': hedef veya gerçekleşen değer okunamadı")
                    eksik = True
                    continue
            h.oran = g
            h.sinirli = min(g, hedef_tavan)
            if g > hedef_tavan:
                uyar("Bilgi", "Hedef tavanı", c, f"'{h.ad}' gerçekleşmesi {yuzde(g)}; {yuzde(hedef_tavan)} ile sınırlandı")
            skor += h.agirlik / toplam_agirlik * h.sinirli
        if eksik:
            continue
        c.skor = skor
        c.bireysel = carpan_bul(skor, skala, kademeli)
        c.carpan = c.bireysel if sirket is None else bireysel_agirlik * c.bireysel + (1 - bireysel_agirlik) * sirket
        if c.hedef_prim is None:
            uyar("Yüksek", "Hedef prim yok", c, "Brüt ücret × maaş katı veya hedef prim tutarı verilmemiş")
            continue
        if c.cikis and c.cikis < bit and not ayrilana_ode:
            c.prim = Decimal(0)
            uyar("Orta", "Dönem içinde ayrıldı", c, f"{c.cikis:%d.%m.%Y} tarihinde ayrıldı; politika gereği ödenmedi (hesaplanan: "
                 f"{tl(c.hedef_prim * c.carpan * c.kist)} TL; --ayrilana-ode)")
            continue
        if c.gun < min_gun:
            c.prim = Decimal(0)
            uyar("Orta", "Asgari çalışma süresi", c, f"Dönemde {c.gun} gün çalıştı (< {min_gun}); prim hakkı yok")
            continue
        c.prim = (c.hedef_prim * c.carpan * c.kist).quantize(K2, ROUND_HALF_UP)
        if c.kist < 1:
            c.notlar.append(f"Kıst {c.gun}/{donem_gun} gün")
        if c.skor < skala[0][0]:
            uyar("Bilgi", "Eşik altı", c, f"Ağırlıklı skor {yuzde(c.skor)} < eşik {yuzde(skala[0][0])}; bireysel çarpan 0")
    return uyarilar


def dagilim(calisanlar: list[Calisan]) -> list[tuple[str, int]]:
    araliklar = [("0 (eşik altı / ödenmedi)", lambda x: x == 0), ("%0–80", lambda x: 0 < x < Decimal("0.8")),
                 ("%80–100", lambda x: Decimal("0.8") <= x < 1), ("%100–120", lambda x: 1 <= x < Decimal("1.2")), ("%120 ve üstü", lambda x: x >= Decimal("1.2"))]
    hesaplanan = [c for c in calisanlar if c.carpan is not None]
    return [(ad, sum(1 for c in hesaplanan if f(c.carpan if c.prim != 0 or c.carpan == 0 else Decimal(0)))) for ad, f in araliklar]


SIRA = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}


def calistir(calisan_yolu: Path, hedef_yolu: Path, skala_yolu: Path, cikti: Path, *, donem: tuple[date, date], butce: Decimal | None = None, **kw) -> dict:
    calisanlar = oku_calisanlar(calisan_yolu)
    skala = oku_skala(skala_yolu)
    uyarilar = hesapla(calisanlar, oku_hedefler(hedef_yolu), skala, donem=donem, **kw)
    toplam = sum((c.prim or Decimal(0) for c in calisanlar), Decimal(0))
    hedef_toplam = sum((c.hedef_prim * c.kist for c in calisanlar if c.hedef_prim is not None and c.prim is not None and (c.prim or not c.carpan)), Decimal(0))
    if butce is not None and toplam > butce:
        uyarilar.append({"onem": "Yüksek", "tur": "Bütçe aşımı", "sicil": "", "aciklama": f"Toplam prim {tl(toplam)} TL > bütçe {tl(butce)} TL. "
                         f"Orantılı ödeme için tüm primler {yuzde(butce / toplam)} ile çarpılabilir (karar yönetimindir)"})
    deps = defaultdict(list)
    for c in calisanlar:
        deps[c.departman].append(c)
    ozet = []
    for d, lst in sorted(deps.items()):
        h = [c for c in lst if c.carpan is not None]
        ozet.append({"departman": d, "n": len(lst), "hesaplanan": len(h), "ort_skor": sum((c.skor for c in h), Decimal(0)) / len(h) if h else None,
                     "ort_carpan": sum((c.carpan for c in h), Decimal(0)) / len(h) if h else None, "prim": sum((c.prim or Decimal(0) for c in lst), Decimal(0))})
        if h and len(h) >= 3 and ozet[-1]["ort_carpan"] >= Decimal("1.2"):
            uyarilar.append({"onem": "Bilgi", "tur": "Kalibrasyon", "sicil": "", "aciklama": f"{d}: ortalama çarpan {yuzde(ozet[-1]['ort_carpan'])}; "
                             "hedeflerin zorluk düzeyi gözden geçirilebilir"})
    uyarilar.sort(key=lambda u: (SIRA[u["onem"]], u["tur"], u["sicil"]))
    s = {"calisanlar": calisanlar, "uyarilar": uyarilar, "toplam": toplam, "hedef_toplam": hedef_toplam, "ozet": ozet, "dagilim": dagilim(calisanlar),
         "skala": skala, "donem": donem, "butce": butce, "ayar": kw}
    rapor_yaz(cikti, s)
    return s


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KONTROL = PatternFill("solid", fgColor="FFF4CE")
UST = Alignment(vertical="top", wrap_text=True)
PF = "#,##0.00"


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def _f(x):
    return None if x is None else float(x)


def rapor_yaz(cikti: Path, s: dict) -> None:
    wb = Workbook()
    pl = wb.active
    pl.title = "Prim Listesi"
    _baslik(pl, ["Sicil", "Ad Soyad", "Departman", "Pozisyon", "Brüt Ücret", "Maaş Katı", "Hedef Prim", "Ağırlıklı Skor", "Bireysel Çarpan", "Toplam Çarpan",
                 "Çalışılan Gün", "Kıst Oranı", "Brüt Prim", "İşaretler", "Onay"], (8, 16, 12, 20, 12, 8, 12, 10, 10, 10, 9, 9, 13, 30, 12))
    for c in s["calisanlar"]:
        pl.append([c.sicil, c.ad, c.departman, c.pozisyon, _f(c.ucret), _f(c.kat), _f(c.hedef_prim), _f(c.skor), _f(c.bireysel), _f(c.carpan), c.gun,
                   float(c.kist), _f(c.prim), "; ".join(c.notlar), ""])
        r = pl.max_row
        for j in (5, 7, 13):
            pl.cell(r, j).number_format = PF
        for j in (8, 9, 10, 12):
            pl.cell(r, j).number_format = "0.0%"
        if any(u["sicil"] == c.sicil and u["onem"] == "Yüksek" for u in s["uyarilar"]):
            pl.cell(r, 14).fill = PatternFill("solid", fgColor="FDE2E1")
        pl.cell(r, 15).fill = KONTROL
    pl.append(["Toplam", "", "", "", None, None, None, None, None, None, None, None, float(s["toplam"])])
    pl.cell(pl.max_row, 13).number_format = PF
    pl.cell(pl.max_row, 13).font = Font(bold=True)
    pl.auto_filter.ref = f"A1:O{pl.max_row - 1}"

    hd = wb.create_sheet("Hedef Detayı")
    _baslik(hd, ["Sicil", "Ad Soyad", "Hedef", "Yön", "Ağırlık", "Hedef Değer", "Gerçekleşen", "Gerçekleşme", "Tavanlı", "Ağırlıklı Katkı"],
            (8, 16, 32, 8, 8, 14, 14, 11, 9, 11))
    for c in s["calisanlar"]:
        top = sum((h.agirlik for h in c.hedefler), Decimal(0)) or Decimal(1)
        for h in c.hedefler:
            hd.append([c.sicil, c.ad, h.ad, "Azalan" if h.yon == "azalan" else "Artan", float(h.agirlik), _f(h.deger), _f(h.gercek), _f(h.oran), _f(h.sinirli),
                       _f(h.agirlik / top * h.sinirli) if h.sinirli is not None else None])
            for j in (8, 9, 10):
                hd.cell(hd.max_row, j).number_format = "0.0%"

    oz = wb.create_sheet("Departman Özeti")
    _baslik(oz, ["Departman", "Çalışan", "Hesaplanan", "Ort. Skor", "Ort. Çarpan", "Toplam Prim"], (18, 9, 11, 10, 10, 14))
    for x in s["ozet"]:
        oz.append([x["departman"], x["n"], x["hesaplanan"], _f(x["ort_skor"]), _f(x["ort_carpan"]), float(x["prim"])])
        for j in (4, 5):
            oz.cell(oz.max_row, j).number_format = "0.0%"
        oz.cell(oz.max_row, 6).number_format = PF
    oz.append([])
    oz.append(["Çarpan dağılımı", "Kişi"])
    for ad, n in s["dagilim"]:
        oz.append([ad, n])
    oz.append([])
    a = s["ayar"]
    d1, d2 = s["donem"]
    bilgiler = [f"Dönem: {d1:%d.%m.%Y} – {d2:%d.%m.%Y}",
                "Skala: " + ", ".join(f"{yuzde(g)} → {yuzde(c)}" for g, c in s["skala"]) + (" (kademeli)" if a.get("kademeli") else " (aralar doğrusal)")
                + "; ilk satırın altı 0, son satırın üstü tavan",
                f"Hedef başına tavan: {yuzde(a.get('hedef_tavan', Decimal('1.5')))}"]
    if a.get("sirket") is not None:
        bilgiler.append(f"Şirket çarpanı: {yuzde(a['sirket'])}, bireysel ağırlık {yuzde(a.get('bireysel_agirlik', Decimal(1)))}")
    bilgiler.append(f"Toplam brüt prim: {tl(s['toplam'])} TL · hedefte (çarpan %100) olsaydı: {tl(s['hedef_toplam'])} TL"
                    + (f" · bütçe: {tl(s['butce'])} TL" if s["butce"] is not None else ""))
    bilgiler.append("Prim ücret niteliğindedir; SGK primi ve gelir vergisi kesintisi bordroda hesaplanır. Kümülatif matrah nedeniyle vergi dilimi "
                    "değişebilir. İşveren maliyeti için SGK işveren payını ekleyin.")
    for b in bilgiler:
        oz.append([b])

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Sicil", "Açıklama", "Karar"], (9, 22, 8, 90, 20))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["sicil"], u["aciklama"], ""])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).alignment = UST
        uy.cell(uy.max_row, 5).fill = KONTROL
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Hedef gerçekleşmesi ve prim skalasından çalışan bazında brüt performans primini hesaplar.")
    p.add_argument("--calisanlar", type=Path, default=ORNEK / "calisanlar.csv", help="Sicil, Ad, Departman, Brüt Ücret, Hedef Prim (Maaş) veya (TL), İşe Giriş, Ayrılış")
    p.add_argument("--hedefler", type=Path, default=ORNEK / "hedefler.csv", help="Sicil, Hedef, Ağırlık (%%), Hedef Değer, Gerçekleşen, Yön (Artan/Azalan)")
    p.add_argument("--skala", type=Path, default=ORNEK / "skala.csv", help="Gerçekleşme (%%), Prim Çarpanı (%%)")
    p.add_argument("--donem", nargs=2, metavar=("BAŞLANGIÇ", "BİTİŞ"), help="Prim dönemi GG.AA.YYYY GG.AA.YYYY (örnek: 2026 1. yarıyıl)")
    p.add_argument("--hedef-tavan", type=float, default=150, help="Hedef başına gerçekleşme tavanı %% (varsayılan 150)")
    p.add_argument("--kademeli", action="store_true", help="Skala satırları arasında doğrusal değil, kademeli çarpan")
    p.add_argument("--sirket-carpani", type=float, help="Şirket performans çarpanı %% (ör. 110)")
    p.add_argument("--bireysel-agirlik", type=float, default=100, help="Şirket çarpanı verildiğinde bireysel çarpanın ağırlığı %% (varsayılan 100)")
    p.add_argument("--ayrilana-ode", action="store_true", help="Dönem bitmeden ayrılana kıst prim öde")
    p.add_argument("--min-gun", type=int, default=0, help="Prim hakkı için dönemde en az çalışılan gün (varsayılan 0)")
    p.add_argument("--butce", help="Toplam prim bütçesi (TL); aşılırsa uyarı ve orantı katsayısı")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "performans_primi.xlsx")
    a = p.parse_args(argv)
    for y in (a.calisanlar, a.hedefler, a.skala):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    if a.donem:
        donem = (tarih(a.donem[0]), tarih(a.donem[1]))
        if None in donem or donem[0] > donem[1]:
            print("[X] --donem iki geçerli tarih olmalı: GG.AA.YYYY GG.AA.YYYY")
            return 2
    elif a.calisanlar == ORNEK / "calisanlar.csv":
        donem = (date(2026, 1, 1), date(2026, 6, 30))
    else:
        print("[X] --donem verin (ör. --donem 01.01.2026 31.12.2026)")
        return 2
    try:
        s = calistir(a.calisanlar, a.hedefler, a.skala, a.cikti, donem=donem, butce=sayi(a.butce) if a.butce else None,
                     hedef_tavan=Decimal(str(a.hedef_tavan)) / YUZ, kademeli=a.kademeli,
                     sirket=Decimal(str(a.sirket_carpani)) / YUZ if a.sirket_carpani is not None else None,
                     bireysel_agirlik=Decimal(str(a.bireysel_agirlik)) / YUZ, ayrilana_ode=a.ayrilana_ode, min_gun=a.min_gun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    n = sum(1 for c in s["calisanlar"] if c.prim)
    print(f"[OK] {len(s['calisanlar'])} çalışan · {n} kişiye prim · toplam brüt {tl(s['toplam'])} TL (hedefte olsaydı {tl(s['hedef_toplam'])} TL)")
    for u in s["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['sicil']} {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
