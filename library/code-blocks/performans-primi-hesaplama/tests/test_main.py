import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
DONEM = (date(2026, 1, 1), date(2026, 6, 30))


def calistir(cikti, **kw):
    return main.calistir(ORNEK / "calisanlar.csv", ORNEK / "hedefler.csv", ORNEK / "skala.csv", cikti, donem=DONEM, **kw)


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "p.xlsx"
        cls.s = calistir(cls.cikti)
        cls.c = {c.sicil: c for c in cls.s["calisanlar"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_skor_ve_carpan(self):
        self.assertEqual((self.c["2001"].skor, self.c["2001"].carpan, self.c["2001"].prim), (D("0.94"), D("0.85"), D("85000.00")))
        self.assertEqual(self.c["2003"].prim, D("338437.50"))      # %107,5 → %118,75 × 285.000
        self.assertEqual(self.c["2004"].skor, D("1.1"))            # azalan fire + hedefi 0 olan iş kazası
        self.assertEqual(self.c["2007"].prim, D("36166.67"))

    def test_kist_ve_ayrilan(self):
        self.assertEqual((self.c["2002"].gun, self.c["2002"].prim), (122, D("72795.58")))   # 108.000 × 122 / 181
        self.assertEqual(self.c["2005"].prim, D("0"))
        self.assertIsNone(self.c["2008"].prim)
        self.assertEqual(self.s["toplam"], D("689768.17"))
        self.assertEqual(self.s["hedef_toplam"].quantize(D("0.01")), D("656707.18"))   # 2005 (ödenmeyen) ve 2008 hariç

    def test_uyarilar_ve_excel(self):
        u = {(x["tur"], x["sicil"]) for x in self.s["uyarilar"]}
        self.assertTrue({("Hedef yok", "2008"), ("Dönem içinde ayrıldı", "2005"), ("Çalışan listesinde yok", "2009")} <= u)
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Prim Listesi", "Hedef Detayı", "Departman Özeti", "Uyarılar"])
        self.assertEqual(wb["Prim Listesi"]["M10"].value, 689768.17)
        self.assertEqual(wb["Hedef Detayı"].max_row, 15)


class KuralTesti(unittest.TestCase):
    def test_skala(self):
        sk = [(D("0.8"), D("0.5")), (D("1"), D("1")), (D("1.2"), D("1.5"))]
        self.assertEqual(main.carpan_bul(D("0.79"), sk), 0)
        self.assertEqual(main.carpan_bul(D("0.9"), sk), D("0.75"))
        self.assertEqual(main.carpan_bul(D("0.9"), sk, kademeli=True), D("0.5"))
        self.assertEqual(main.carpan_bul(D("2"), sk), D("1.5"))

    def test_gerceklesme(self):
        H = main.Hedef
        self.assertEqual(main.gerceklesme(H("x", D(1), D(0), D(2), "azalan")), 0)
        self.assertEqual(main.gerceklesme(H("x", D(1), D(4), D(5), "azalan")), D("0.8"))
        self.assertIsNone(main.gerceklesme(H("x", D(1), None, D(5), "artan")))

    def test_secenekler(self):
        with tempfile.TemporaryDirectory() as tmp:
            s = calistir(Path(tmp) / "p.xlsx", sirket=D("1.1"), bireysel_agirlik=D("0.7"), ayrilana_ode=True, butce=D("700000"), hedef_tavan=D("1.2"))
            c = {x.sicil: x for x in s["calisanlar"]}
            self.assertEqual(c["2001"].prim, D("92500.00"))               # 0,7 × 0,85 + 0,3 × 1,10
            self.assertEqual(c["2005"].skor, D("1.1"))                     # şikâyet %133 → %120 tavan
            self.assertGreater(c["2005"].prim, 0)
            turler = {u["tur"] for u in s["uyarilar"]}
            self.assertTrue({"Bütçe aşımı", "Hedef tavanı"} <= turler)

    def test_agirlik_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t / "c.csv").write_text("Sicil No;Brüt Ücret;Hedef Prim (TL);İşe Giriş\n1;40000;10000;01.01.2020\n2;40000;10000;15.06.2026\n", encoding="utf-8")
            (t / "h.csv").write_text("Sicil No;Hedef;Ağırlık;Gerçekleşme (%)\n1;A;30;100\n1;B;30;120\n2;A;100;100\n", encoding="utf-8")
            s = main.calistir(t / "c.csv", t / "h.csv", ORNEK / "skala.csv", t / "o.xlsx", donem=DONEM, min_gun=30)
            c = {x.sicil: x for x in s["calisanlar"]}
            self.assertEqual((c["1"].skor, c["1"].prim), (D("1.1"), D("12500.00")))
            self.assertEqual(c["2"].prim, 0)
            self.assertTrue(any(u["tur"] == "Ağırlık toplamı" for u in s["uyarilar"]))
            self.assertEqual(main.main(["--cikti", str(t / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--calisanlar", str(t / "c.csv"), "--cikti", str(t / "y.xlsx")]), 2)
            self.assertEqual(main.main(["--skala", str(t / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
