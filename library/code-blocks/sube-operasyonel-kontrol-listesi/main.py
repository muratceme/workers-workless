"""
Şube Operasyonel Kontrol Listesi — Workers / Workless kod bloğu
Şube Bankacılığı › Şube Operasyon Yöneticisi

Günlük, haftalık ve aylık periyodik kontrolleri (kasa, kıymetli evrak, anahtar, yetki) sorumlu ve tamamlanma durumuyla izler.
  Sonraki tarih = Günlük   : son yapılma tarihinden sonraki ilk iş günü
                  Haftalık : son yapılma + 7 gün · Aylık: + 1 ay · Üç Aylık: + 3 ay · Altı Aylık: + 6 ay · Yıllık: + 12 ay
                  (hafta sonuna ya da tatil listesindeki bir güne gelirse sonraki ilk iş gününe kayar)
  Durum         = Hiç yapılmamış (son yapılma tarihi yok) · Gecikti (durum tarihi > sonraki tarih) · Bugün · Tamam
  Uyum %        = (kontrol sayısı − geciken − hiç yapılmamış) ÷ kontrol sayısı × 100
Kontrollerin neler olacağını ve sıklığını bankanızın iç düzenlemeleri belirler; kod yalnız yazdığınız listeyi izler.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
import calendar

SUTUN = {"ad": ("kontrol", "kontrol adi", "kontrol noktasi"), "periyot": ("periyot", "siklik"), "sorumlu": ("sorumlu", "kontrol eden", "gorevli"), "son": ("son yapilma tarihi", "son yapilma", "son kontrol tarihi", "son kontrol"),
         "sonuc": ("sonuc", "son sonuc"), "nota": ("not", "aciklama")}
PERIYOT = {"gunluk": ("Günlük", 0), "her gun": ("Günlük", 0), "haftalik": ("Haftalık", 0), "aylik": ("Aylık", 1), "uc aylik": ("Üç Aylık", 3), "3 aylik": ("Üç Aylık", 3), "ceyreklik": ("Üç Aylık", 3), "alti aylik": ("Altı Aylık", 6),
           "6 aylik": ("Altı Aylık", 6), "yillik": ("Yıllık", 12)}
P_SIRA = ("Günlük", "Haftalık", "Aylık", "Üç Aylık", "Altı Aylık", "Yıllık")
BULGU = {"bulgu", "bulgu var", "uygun degil", "uygunsuz", "eksik", "fark var"}
DURUMLAR = ("Hiç yapılmamış", "Gecikti", "Bugün", "Tamam")
DURUM_RENK = {"Hiç yapılmamış": KIRMIZI, "Gecikti": KIRMIZI, "Bugün": SARI, "Tamam": YESIL}


def ay_ekle(t: date, n: int) -> date:
    y, a = divmod(t.year * 12 + t.month - 1 + n, 12)
    return date(y, a + 1, min(t.day, calendar.monthrange(y, a + 1)[1]))


def is_gunu(t: date, tatiller: set[date]) -> date:
    """t iş günüyse kendisi, değilse sonraki ilk iş günü."""
    while t.weekday() >= 5 or t in tatiller:
        t += timedelta(days=1)
    return t


def sonraki(son: date, periyot: str, ay: int, tatiller: set[date]) -> date:
    if periyot == "Günlük":
        return is_gunu(son + timedelta(days=1), tatiller)
    return is_gunu(son + timedelta(days=7) if periyot == "Haftalık" else ay_ekle(son, ay), tatiller)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("ad", "periyot")):
        ad = metin(r["ad"])
        if not ad:
            continue
        p = PERIYOT.get(katla(r["periyot"]))
        if p is None:
            uy.append(uyari("Yüksek", "Periyot okunamadı", ad, f"Satır {r['_satir']}: periyot '{metin(r['periyot'])}'. Günlük, Haftalık, Aylık, Üç Aylık, Altı Aylık ya da Yıllık yazın. Satır alınmadı."))
            continue
        if metin(r["son"]) and tarih(r["son"]) is None:
            uy.append(uyari("Orta", "Tarih okunamadı", ad, f"Son yapılma tarihi '{metin(r['son'])}' okunamadı; boş sayıldı."))
        l.append({"ad": ad, "periyot": p[0], "ay": p[1], "sorumlu": metin(r["sorumlu"]) or "(sorumlu yok)", "son": tarih(r["son"]), "sonuc": metin(r["sonuc"]), "bulgu": katla(r["sonuc"]) in BULGU, "nota": metin(r["nota"])})
    return l, uy


def oku_tatiller(yol: Path) -> set[date]:
    return {t for r in tablo_oku(yol) for t in [tarih(r[0])] if t}


def degerlendir(kontroller: list[dict], bugun: date, tatiller: set[date] | None = None) -> dict:
    if not kontroller:
        raise ValueError("İzlenecek kontrol yok.")
    tatiller, uy = tatiller or set(), []
    for k in kontroller:
        k["sonraki"] = sonraki(k["son"], k["periyot"], k["ay"], tatiller) if k["son"] else None
        k["gecikme"] = max((bugun - k["sonraki"]).days, 0) if k["sonraki"] else None
        kim = f"{k['ad']} · {k['sorumlu']}"
        if k["son"] is None:
            k["durum"] = "Hiç yapılmamış"
            uy.append(uyari("Yüksek", "Hiç yapılmamış", kim, f"{k['periyot']} kontrol için son yapılma tarihi yok."))
        elif bugun > k["sonraki"]:
            k["durum"] = "Gecikti"
            uy.append(uyari("Yüksek", "Kontrol gecikti", kim, f"{k['periyot']} kontrol en son {k['son']:%d.%m.%Y} tarihinde yapılmış; {k['sonraki']:%d.%m.%Y} tarihinde yapılmalıydı ({k['gecikme']} gün geçti)."))
        elif bugun == k["sonraki"]:
            k["durum"] = "Bugün"
            uy.append(uyari("Bilgi", "Bugün yapılacak", kim, f"{k['periyot']} kontrol bugün yapılmalı (son yapılma {k['son']:%d.%m.%Y})."))
        else:
            k["durum"] = "Tamam"
        if k["bulgu"]:
            uy.append(uyari("Yüksek", "Son kontrolde bulgu var", kim, f"Son kontrolün sonucu '{k['sonuc']}'" + (f": {k['nota']}" if k["nota"] else "") + ". Bulgunun giderildiğini teyit edin."))
        if k["son"] and k["son"] > bugun:
            uy.append(uyari("Orta", "Tarih ileride", kim, f"Son yapılma tarihi {k['son']:%d.%m.%Y}, durum tarihi {bugun:%d.%m.%Y}."))

    def ozet(alan: str, sira=None) -> list[dict]:
        grup = {}
        for k in kontroller:
            g = grup.setdefault(katla(k[alan]), {"ad": k[alan], "kontrol": 0, **{d: 0 for d in DURUMLAR}, "bulgu": 0})
            g["kontrol"] += 1
            g[k["durum"]] += 1
            g["bulgu"] += k["bulgu"]
        return sorted(grup.values(), key=lambda g: sira.index(g["ad"]) if sira else katla(g["ad"]))

    say = {d: sum(1 for k in kontroller if k["durum"] == d) for d in DURUMLAR}
    isler = sorted((k for k in kontroller if k["durum"] != "Tamam"), key=lambda k: (DURUMLAR.index(k["durum"]), k["sonraki"] or date.min, katla(k["ad"])))
    return {"kontroller": kontroller, "isler": isler, "say": say, "bulgu": sum(1 for k in kontroller if k["bulgu"]), "uyum": yuvarla(Decimal((len(kontroller) - say["Gecikti"] - say["Hiç yapılmamış"]) * 100) / len(kontroller)),
            "periyotlar": ozet("periyot", P_SIRA), "sorumlular": ozet("sorumlu"), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, bugun: date) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Bugünün Listesi"
    _baslik(ws, ["Durum", "Kontrol", "Periyot", "Sorumlu", "Son Yapılma", "Yapılması Gereken", "Gecikme (gün)", "Yapıldı (tarih / paraf)", "Sonuç / Not"], (18, 46, 12, 28, 12, 13, 10, 22, 34))
    for k in r["isler"]:
        _satir(ws, [k["durum"], k["ad"], k["periyot"], k["sorumlu"], k["son"], k["sonraki"], k["gecikme"] or None, None, None], tarih_s=(5, 6), dolgu={1: DURUM_RENK[k["durum"]]}, sar=(2,))
    ws.append([])
    ws.append([f"Durum tarihi: {bugun:%d.%m.%Y} · uyum %{r['uyum']}"])
    ws = wb.create_sheet("Tüm Kontroller")
    _baslik(ws, ["Durum", "Kontrol", "Periyot", "Sorumlu", "Son Yapılma", "Sonraki Tarih", "Son Sonuç", "Not"], (18, 46, 12, 28, 12, 13, 14, 40))
    for k in sorted(r["kontroller"], key=lambda k: (P_SIRA.index(k["periyot"]), katla(k["ad"]))):
        _satir(ws, [k["durum"], k["ad"], k["periyot"], k["sorumlu"], k["son"], k["sonraki"], k["sonuc"], k["nota"]], tarih_s=(5, 6), dolgu={1: DURUM_RENK[k["durum"]], 7: KIRMIZI if k["bulgu"] else None}, sar=(2, 8))
    for baslik, alan, ad in (("Periyot Özeti", "periyotlar", "Periyot"), ("Sorumlu Özeti", "sorumlular", "Sorumlu")):
        ws = wb.create_sheet(baslik)
        _baslik(ws, [ad, "Kontrol"] + list(DURUMLAR) + ["Bulgulu"], (30, 9, 13, 9, 9, 9, 9))
        for g in r[alan]:
            _satir(ws, [g["ad"], g["kontrol"]] + [g[d] or None for d in DURUMLAR] + [g["bulgu"] or None], dolgu={3: KIRMIZI if g["Hiç yapılmamış"] else None, 4: KIRMIZI if g["Gecikti"] else None})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, tatil_yol: Path | None = None) -> dict:
    kontroller, uy = oku(yol)
    r = degerlendir(kontroller, bugun, oku_tatiller(tatil_yol) if tatil_yol else None)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, bugun)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Şubenin günlük, haftalık ve aylık periyodik kontrollerini sorumlu ve tamamlanma durumuyla izler.")
    p.add_argument("--kontroller", type=Path, help="Kontrol, Periyot, Sorumlu, Son Yapılma Tarihi, Sonuç, Not")
    p.add_argument("--tatiller", type=Path, help="Resmî tatil tarihleri, her satırda bir tarih (isteğe bağlı)")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "sube_operasyonel_kontrol_listesi.xlsx")
    a = p.parse_args(argv)
    ornek = a.kontroller is None
    yol = ORNEK / "kontroller.csv" if ornek else a.kontroller
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    if bugun is None:
        print("[X] --tarih GG.AA.YYYY biçiminde olmalı")
        return 2
    for y in (yol, a.tatiller):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(yol, a.cikti, bugun, a.tatiller)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    s = r["say"]
    print(f"[OK] {len(r['kontroller'])} kontrol · gecikti {s['Gecikti']}, hiç yapılmamış {s['Hiç yapılmamış']}, bugün {s['Bugün']}, tamam {s['Tamam']} · uyum %{r['uyum']}" + (f" · bulgulu {r['bulgu']}" if r["bulgu"] else ""))
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
