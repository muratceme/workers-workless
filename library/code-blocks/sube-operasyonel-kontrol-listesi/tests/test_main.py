import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)   # Pazartesi
AY = {"Günlük": 0, "Haftalık": 0, "Aylık": 1, "Üç Aylık": 3, "Altı Aylık": 6, "Yıllık": 12}


def k(ad, son, periyot="Günlük", sorumlu="S", sonuc="Uygun"):
    return {"ad": ad, "periyot": periyot, "ay": AY[periyot], "sorumlu": sorumlu, "son": son, "sonuc": sonuc, "bulgu": sonuc == "Bulgu", "nota": ""}


class TarihTesti(unittest.TestCase):
    def test_sonraki(self):
        t = {date(2026, 10, 29), date(2026, 10, 13)}
        s = lambda son, p, tatil=frozenset(): main.sonraki(date(*son), p, AY[p], set(tatil))   # noqa: E731
        self.assertEqual([s((2026, 10, 8), "Günlük"), s((2026, 10, 9), "Günlük"), s((2026, 10, 10), "Günlük"), s((2026, 10, 12), "Günlük", t), s((2026, 10, 28), "Günlük", t)],
                         [date(2026, 10, 9), date(2026, 10, 12), date(2026, 10, 12), date(2026, 10, 14), date(2026, 10, 30)])
        self.assertEqual([s((2026, 10, 5), "Haftalık"), s((2026, 10, 3), "Haftalık"), s((2026, 10, 22), "Haftalık", t), s((2026, 9, 15), "Aylık"), s((2026, 1, 31), "Aylık"), s((2026, 7, 10), "Üç Aylık"), s((2026, 4, 12), "Altı Aylık"),
                          s((2025, 10, 11), "Yıllık")],
                         [date(2026, 10, 12), date(2026, 10, 12), date(2026, 10, 30), date(2026, 10, 15), date(2026, 3, 2), date(2026, 10, 12), date(2026, 10, 12), date(2026, 10, 12)])
        self.assertEqual([main.is_gunu(date(2026, 10, d), {date(2026, 10, 12)}) for d in (9, 10, 11, 12)], [date(2026, 10, 9), date(2026, 10, 13), date(2026, 10, 13), date(2026, 10, 13)])


class DegerlendirTesti(unittest.TestCase):
    def test_durumlar(self):
        l = [k("cuma", date(2026, 10, 9)), k("perşembe", date(2026, 10, 8)), k("bugün", BUGUN), k("hiç", None), k("bulgu", date(2026, 10, 9), sonuc="Bulgu"), k("ileri", date(2026, 10, 20)), k("hafta", date(2026, 10, 6), "Haftalık"),
             k("ay", date(2026, 9, 11), "Aylık")]
        r = main.degerlendir(l, BUGUN)
        self.assertEqual([(x["ad"], x["durum"], x["sonraki"], x["gecikme"]) for x in r["kontroller"]],
                         [("cuma", "Bugün", BUGUN, 0), ("perşembe", "Gecikti", date(2026, 10, 9), 3), ("bugün", "Tamam", date(2026, 10, 13), 0), ("hiç", "Hiç yapılmamış", None, None), ("bulgu", "Bugün", BUGUN, 0),
                          ("ileri", "Tamam", date(2026, 10, 21), 0), ("hafta", "Tamam", date(2026, 10, 13), 0), ("ay", "Bugün", date(2026, 10, 12), 0)])
        self.assertEqual((r["say"], r["bulgu"], r["uyum"], [x["ad"] for x in r["isler"]]), ({"Hiç yapılmamış": 1, "Gecikti": 1, "Bugün": 3, "Tamam": 3}, 1, D("75.00"), ["hiç", "perşembe", "ay", "bulgu", "cuma"]))
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         sorted([("Bilgi", "Bugün yapılacak", "cuma · S"), ("Yüksek", "Kontrol gecikti", "perşembe · S"), ("Yüksek", "Hiç yapılmamış", "hiç · S"), ("Bilgi", "Bugün yapılacak", "bulgu · S"),
                                 ("Yüksek", "Son kontrolde bulgu var", "bulgu · S"), ("Orta", "Tarih ileride", "ileri · S"), ("Bilgi", "Bugün yapılacak", "ay · S")]))
        self.assertEqual(main.degerlendir([k("a", date(2026, 10, 9))], BUGUN, {BUGUN})["kontroller"][0]["durum"], "Tamam")
        self.assertEqual(main.degerlendir([k("a", date(2026, 10, 9))], date(2026, 10, 11))["kontroller"][0]["durum"], "Tamam")
        with self.assertRaises(ValueError):
            main.degerlendir([], BUGUN)

    def test_ozet(self):
        l = [k("1", date(2026, 10, 9), "Günlük", "Ali"), k("2", None, "Aylık", "ali"), k("3", date(2026, 10, 1), "Haftalık", "Ece", "Bulgu"), k("4", date(2026, 10, 12), "Günlük", "Ece")]
        r = main.degerlendir(l, BUGUN)
        o = lambda alan: [(g["ad"], g["kontrol"], g["Hiç yapılmamış"], g["Gecikti"], g["Bugün"], g["Tamam"], g["bulgu"]) for g in r[alan]]   # noqa: E731
        self.assertEqual((o("periyotlar"), o("sorumlular"), r["uyum"]), ([("Günlük", 2, 0, 0, 1, 1, 0), ("Haftalık", 1, 0, 1, 0, 0, 1), ("Aylık", 1, 1, 0, 0, 0, 0)], [("Ali", 2, 1, 0, 1, 0, 0), ("Ece", 2, 0, 1, 0, 1, 1)], D("50.00")))


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "kontroller.csv", cikti, BUGUN)
            self.assertEqual([(x["durum"], x["sonraki"], x["gecikme"]) for x in r["kontroller"]],
                             [("Bugün", BUGUN, 0), ("Gecikti", date(2026, 10, 9), 3), ("Tamam", date(2026, 10, 13), 0), ("Bugün", BUGUN, 0), ("Bugün", BUGUN, 0), ("Gecikti", date(2026, 10, 5), 7), ("Tamam", date(2026, 10, 15), 0),
                              ("Gecikti", date(2026, 10, 1), 11), ("Hiç yapılmamış", None, None), ("Bugün", BUGUN, 0), ("Tamam", date(2026, 10, 20), 0)])
            self.assertEqual((r["say"], r["bulgu"], r["uyum"]), ({"Hiç yapılmamış": 1, "Gecikti": 3, "Bugün": 4, "Tamam": 3}, 1, D("63.64")))
            self.assertEqual([x["ad"] for x in r["isler"]][:4], ["Kiralık kasa ziyaret defteri kontrolü", "Bekleyen fişler ve geçici hesap kontrolü", "Alarm ve panik butonu testi", "ATM kasa mutabakatı"])
            o = lambda alan: [(g["ad"], g["kontrol"], g["Hiç yapılmamış"], g["Gecikti"], g["Bugün"], g["Tamam"]) for g in r[alan]]   # noqa: E731
            self.assertEqual(o("periyotlar"), [("Günlük", 4, 0, 1, 2, 1), ("Haftalık", 2, 0, 1, 1, 0), ("Aylık", 4, 1, 1, 0, 2), ("Üç Aylık", 1, 0, 0, 1, 0)])
            self.assertEqual(o("sorumlular"), [("Güvenlik Görevlisi", 2, 0, 1, 0, 1), ("Operasyon Yetkilisi", 3, 1, 1, 1, 0), ("Şube Müdürü", 2, 0, 0, 1, 1), ("Şube Operasyon Yöneticisi", 4, 0, 1, 2, 1)])
            self.assertEqual([(u["onem"], u["tur"]) for u in r["uyarilar"]], [("Yüksek", "Hiç yapılmamış")] + [("Yüksek", "Kontrol gecikti")] * 3 + [("Yüksek", "Son kontrolde bulgu var")] + [("Bilgi", "Bugün yapılacak")] * 4)
            self.assertEqual(load_workbook(cikti).sheetnames, ["Bugünün Listesi", "Tüm Kontroller", "Periyot Özeti", "Sorumlu Özeti", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, ty = Path(tmp) / "k.csv", Path(tmp) / "t.csv"
            y.write_text("Kontrol Adı;Sıklık;Görevli;Son Kontrol;Sonuç\nKasa;HER GÜN;;09.10.2026;UYGUN DEĞİL\nAlarm;3 Aylık;Ali;x;\nYetki;iki haftada bir;;;\n;Günlük;;;\n", encoding="utf-8")
            ty.write_text("Tarih;Ad\n29.10.2026;Cumhuriyet Bayramı\nx;y\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["ad"], x["periyot"], x["ay"], x["sorumlu"], x["son"], x["bulgu"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("Kasa", "Günlük", 0, "(sorumlu yok)", date(2026, 10, 9), True), ("Alarm", "Üç Aylık", 3, "Ali", None, False)], [("Tarih okunamadı", "Alarm"), ("Periyot okunamadı", "Yetki")]))
            self.assertEqual(main.oku_tatiller(ty), {date(2026, 10, 29)})
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--kontroller", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tatiller", str(ty), "--tarih", "30.10.2026"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(arg + ["--tarih", "x"]), 2)
            self.assertEqual(main.main(arg + ["--tatiller", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Kontrol;Sorumlu\nA;B\n", "Kontrol;Periyot\n;Günlük\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
