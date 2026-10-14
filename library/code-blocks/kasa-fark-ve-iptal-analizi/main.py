"""
Kasa Fark ve İptal Analizi — Workers / Workless kod bloğu
Mağaza Operasyonları › Kasa Sorumlusu

Kasiyer bazında iade, iptal ve kasa farklarını dönem boyunca toplar; diğer kasiyerlerden belirgin ayrışan örüntüleri işaretler.
  İade % / iptal %        = iade ya da iptal tutarı ÷ satış tutarı × 100;   100 fişte iade / iptal = adet ÷ fiş adedi × 100
  Farklı gün              = kasa farkının mutlak değeri --tolerans (5) üstünde olan gün;   eksik gün: fark eksi, fazla gün: fark artı
  En uzun eksik serisi    = kasiyerin çalıştığı günler sırasıyla, art arda eksik veren gün sayısı
  İşaretler (fiş adedi --min, 100 ve üstü olan kasiyerde):
      İade oranı yüksek   = kasiyerin iade %'si, diğer kasiyerlerin toplam iade %'sinin --kat (2) katından büyük
      İptal oranı yüksek  = aynı karşılaştırma iptal için
      Sürekli kasa eksiği = en uzun eksik serisi --seri (3) gün ve üstü
      Sık kasa farkı      = farklı gün ÷ çalışılan gün × 100 değeri --fark-gun (%50) üstünde
İşaret bir inceleme önerisidir, suistimal tespiti değildir; vardiya, kasa konumu ve ürün karması oranları doğal olarak değiştirir.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"tarih": ("tarih", "gun"), "kasiyer": ("kasiyer", "personel", "kasa sorumlusu"), "fis": ("fis adedi", "fis sayisi", "islem adedi"), "satis": ("satis tutari", "satis", "ciro", "net satis"),
         "iade_adet": ("iade adedi", "iade sayisi"), "iade": ("iade tutari", "iade"), "iptal_adet": ("iptal adedi", "iptal sayisi"), "iptal": ("iptal tutari", "iptal"), "fark": ("kasa farki", "fark", "nakit farki")}
SAYI = ("fis", "satis", "iade_adet", "iade", "iptal_adet", "iptal")
SIFIR = Decimal(0)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("tarih", "kasiyer", "satis")):
        kasiyer = metin(r["kasiyer"])
        if not kasiyer:
            continue
        t, d, fark = tarih(r["tarih"]), {a: para(r[a]) if metin(r[a]) else SIFIR for a in SAYI}, para(r["fark"]) if metin(r["fark"]) else SIFIR
        if t is None or fark is None or any(v is None or v < 0 for v in d.values()):
            uy.append(uyari("Yüksek", "Satır okunamadı", f"{kasiyer} · {metin(r['tarih'])}", f"Satır {r['_satir']}: tarih ya da sayılardan biri okunamadı. Satır alınmadı."))
            continue
        l.append({"tarih": t, "kasiyer": kasiyer, **d, "fark": fark})
    return l, uy


def _oran(pay, payda) -> Decimal | None:
    return yuvarla(Decimal(pay) * 100 / Decimal(payda)) if payda else None


def gosterge(l: list[dict], tolerans: Decimal) -> dict:
    t = {a: sum((g[a] for g in l), SIFIR) for a in SAYI}
    seri = enb = 0
    for g in sorted(l, key=lambda g: g["tarih"]):
        seri = seri + 1 if g["fark"] < -tolerans else 0
        enb = max(enb, seri)
    farkli = sum(1 for g in l if abs(g["fark"]) > tolerans)
    return {**t, "gun": len(l), "iade_oran": _oran(t["iade"], t["satis"]), "iptal_oran": _oran(t["iptal"], t["satis"]), "iade_100": _oran(t["iade_adet"], t["fis"]), "iptal_100": _oran(t["iptal_adet"], t["fis"]),
            "fark": sum((g["fark"] for g in l), SIFIR), "mutlak": sum((abs(g["fark"]) for g in l), SIFIR), "farkli": farkli, "eksik": sum(1 for g in l if g["fark"] < -tolerans), "fazla": sum(1 for g in l if g["fark"] > tolerans),
            "farkli_oran": _oran(farkli, len(l)), "seri": enb}


def analiz(gunler: list[dict], tolerans: Decimal = Decimal(5), kat: Decimal = Decimal(2), seri: int = 3, fark_gun: Decimal = Decimal(50), asgari: int = 100) -> dict:
    if not gunler:
        raise ValueError("Analiz edilecek kasa kaydı yok.")
    grup, uy = {}, []
    for g in gunler:
        g["farkli"] = abs(g["fark"]) > tolerans
        grup.setdefault(katla(g["kasiyer"]), []).append(g)
    genel = gosterge(gunler, tolerans)
    kasiyerler = []
    for anahtar, g in grup.items():
        k = {"kasiyer": g[0]["kasiyer"], **gosterge(g, tolerans)}
        diger = gosterge([x for a, l in grup.items() if a != anahtar for x in l], tolerans) if len(grup) > 1 else None
        k["diger_iade"], k["diger_iptal"] = (diger["iade_oran"], diger["iptal_oran"]) if diger else (None, None)
        isaret = []
        if k["fis"] >= asgari:
            if k["iade_oran"] is not None and k["diger_iade"] and k["iade_oran"] > k["diger_iade"] * kat:
                isaret.append("İade oranı yüksek")
            if k["iptal_oran"] is not None and k["diger_iptal"] and k["iptal_oran"] > k["diger_iptal"] * kat:
                isaret.append("İptal oranı yüksek")
            if k["seri"] >= seri:
                isaret.append("Sürekli kasa eksiği")
            if k["farkli_oran"] > fark_gun:
                isaret.append("Sık kasa farkı")
        k["isaretler"], k["durum"] = isaret, "Az veri" if k["fis"] < asgari else "; ".join(isaret) or "Olağan"
        if isaret:
            uy.append(uyari("Yüksek" if len(isaret) >= 2 else "Orta", "Kasiyer diğerlerinden ayrışıyor", k["kasiyer"], f"{k['durum']}. İade %{k['iade_oran']} (diğerleri %{k['diger_iade']}), iptal %{k['iptal_oran']} (diğerleri %{k['diger_iptal']}); "
                            f"{k['gun']} günün {k['farkli']} tanesinde kasa farkı, en uzun eksik serisi {k['seri']} gün, net fark {k['fark']:+,.2f}. Fiş ve kamera kayıtlarıyla inceleyin."))
        kasiyerler.append(k)
    kasiyerler.sort(key=lambda k: (-len(k["isaretler"]), k["durum"] == "Az veri", -(k["iade_oran"] or 0) - (k["iptal_oran"] or 0), katla(k["kasiyer"])))
    return {"gunler": sorted(gunler, key=lambda g: (katla(g["kasiyer"]), g["tarih"])), "kasiyerler": kasiyerler, "genel": genel, "uyarilar": uy}


BASLIK = ["Gün", "Fiş", "Satış", "İade Adedi", "İade Tutarı", "İade %", "100 Fişte İade", "İptal Adedi", "İptal Tutarı", "İptal %", "100 Fişte İptal", "Net Kasa Farkı", "Mutlak Fark Toplamı", "Farklı Gün", "Eksik Gün", "Fazla Gün",
          "En Uzun Eksik Serisi"]


def _deger(k: dict) -> list:
    return [k["gun"], k["fis"], k["satis"], k["iade_adet"], k["iade"], k["iade_oran"], k["iade_100"], k["iptal_adet"], k["iptal"], k["iptal_oran"], k["iptal_100"], k["fark"], k["mutlak"], k["farkli"] or None, k["eksik"] or None, k["fazla"] or None,
            k["seri"] or None]


def rapor_yaz(cikti: Path, r: dict, tolerans: Decimal) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kasiyerler"
    _baslik(ws, ["Kasiyer", "Değerlendirme"] + BASLIK, (22, 44, 6, 8, 15, 8, 13, 8, 9, 8, 13, 8, 9, 13, 13, 8, 8, 8, 10), "B2")
    for k in r["kasiyerler"]:
        _satir(ws, [k["kasiyer"], k["durum"]] + _deger(k), para_s=(5, 7, 8, 9, 11, 12, 13, 14, 15),
               dolgu={2: None if k["durum"] == "Az veri" else YESIL if k["durum"] == "Olağan" else KIRMIZI if len(k["isaretler"]) >= 2 else SARI, 8: SARI if "İade oranı yüksek" in k["isaretler"] else None,
                      12: SARI if "İptal oranı yüksek" in k["isaretler"] else None, 19: SARI if "Sürekli kasa eksiği" in k["isaretler"] else None}, sar=(2,))
    n = _satir(ws, ["MAĞAZA", None] + _deger(r["genel"]), para_s=(5, 7, 8, 9, 11, 12, 13, 14, 15))
    for h in ws[n]:
        h.font = Font(bold=True)
    ws.append([])
    ws.append(["Değerlendirme bir inceleme önerisidir; suistimal tespiti değildir. Oranlar diğer kasiyerlerin toplam oranıyla karşılaştırılır."])
    ws = wb.create_sheet("Günlük Kayıtlar")
    _baslik(ws, ["Kasiyer", "Tarih", "Fiş", "Satış", "İade Adedi", "İade Tutarı", "İade %", "İptal Adedi", "İptal Tutarı", "İptal %", "Kasa Farkı"], (22, 12, 8, 15, 8, 13, 8, 8, 13, 8, 12))
    for g in r["gunler"]:
        _satir(ws, [g["kasiyer"], g["tarih"], g["fis"], g["satis"], g["iade_adet"] or None, g["iade"] or None, _oran(g["iade"], g["satis"]), g["iptal_adet"] or None, g["iptal"] or None, _oran(g["iptal"], g["satis"]), g["fark"]],
               para_s=(4, 6, 7, 9, 10, 11), tarih_s=(2,), dolgu={11: None if not g["farkli"] else KIRMIZI if g["fark"] < 0 else SARI})
    ws.append([])
    ws.append([f"Kasa farkı toleransı {tolerans}; kırmızı: eksik, sarı: fazla."])
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, tolerans: Decimal = Decimal(5), kat: Decimal = Decimal(2), seri: int = 3, fark_gun: Decimal = Decimal(50), asgari: int = 100) -> dict:
    gunler, uy = oku(yol)
    r = analiz(gunler, tolerans, kat, seri, fark_gun, asgari)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, tolerans)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Kasiyer bazında iade, iptal ve kasa farklarını analiz edip olağandışı örüntüleri işaretler.")
    p.add_argument("--kayitlar", type=Path, help="Tarih, Kasiyer, Fiş Adedi, Satış Tutarı, İade Adedi, İade Tutarı, İptal Adedi, İptal Tutarı, Kasa Farkı")
    p.add_argument("--tolerans", default="5", help="Bu tutara kadar kasa farkı yok sayılır (varsayılan 5)")
    p.add_argument("--kat", default="2", help="Kasiyerin oranı diğerlerinin oranının kaç katını aşarsa işaretlensin (varsayılan 2)")
    p.add_argument("--seri", type=int, default=3, help="Art arda kaç gün eksik 'sürekli kasa eksiği' sayılsın (varsayılan 3)")
    p.add_argument("--fark-gun", default="50", help="Farklı günlerin payı bu yüzdeyi aşarsa 'sık kasa farkı' (varsayılan 50)")
    p.add_argument("--min", type=int, default=100, dest="asgari", help="Değerlendirme için dönemde en az fiş adedi (varsayılan 100)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kasa_fark_iptal_analizi.xlsx")
    a = p.parse_args(argv)
    yol = a.kayitlar or ORNEK / "kasiyer_gunluk.csv"
    tol, kat, fg = para(a.tolerans), para(a.kat), para(a.fark_gun)
    if tol is None or tol < 0 or kat is None or kat < 1 or a.seri < 2 or fg is None or not 0 <= fg <= 100 or a.asgari < 1:
        print("[X] --tolerans eksi olamaz; --kat en az 1; --seri en az 2; --fark-gun 0 ile 100 arasında; --min en az 1 olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, tol, kat, a.seri, fg, a.asgari)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    g = r["genel"]
    print(f"[OK] {len(r['kasiyerler'])} kasiyer · {g['gun']} kasiyer-gün · satış {g['satis']:,.2f} · iade %{g['iade_oran']} · iptal %{g['iptal_oran']} · net kasa farkı {g['fark']:+,.2f} · işaretlenen {sum(1 for k in r['kasiyerler'] if k['isaretler'])}")
    for k in r["kasiyerler"]:
        if k["isaretler"]:
            print(f"[!] {k['kasiyer']}: {k['durum']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
