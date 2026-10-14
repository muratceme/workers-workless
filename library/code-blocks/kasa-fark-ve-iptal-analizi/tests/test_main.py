import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def g(kasiyer, gun, fark, iade=100, iptal=50, fis=100, satis=10000, iade_adet=2, iptal_adet=1):
    return {"tarih": date(2026, 10, gun), "kasiyer": kasiyer, "fis": D(fis), "satis": D(satis), "iade_adet": D(iade_adet), "iade": D(iade), "iptal_adet": D(iptal_adet), "iptal": D(iptal), "fark": D(str(fark))}


def veri():
    return [g("A", 2, -20), g("a", 1, -10), g("A", 3, -6), g("A", 4, 0)] + [g("B", i + 1, f, 800) for i, f in enumerate((0, 0, 3, -2))] + [g("C", 1, -100, 1000, 500, 50, 5000)]


class AnalizTesti(unittest.TestCase):
    def test_gostergeler_ve_isaretler(self):
        r = main.analiz(veri())
        self.assertEqual([(k["kasiyer"], k["durum"], k["gun"], k["fis"], k["satis"], k["iade"], k["iade_oran"], k["iptal_oran"], k["iade_100"], k["fark"], k["mutlak"], k["farkli"], k["eksik"], k["fazla"], k["farkli_oran"], k["seri"], k["diger_iade"],
                           k["diger_iptal"]) for k in r["kasiyerler"]],
                         [("A", "Sürekli kasa eksiği; Sık kasa farkı", 4, D(400), D(40000), D(400), D("1.00"), D("0.50"), D("2.00"), D(-36), D(36), 3, 3, 0, D("75.00"), 3, D("9.33"), D("1.56")),
                          ("B", "İade oranı yüksek", 4, D(400), D(40000), D(3200), D("8.00"), D("0.50"), D("2.00"), D(1), D(5), 0, 0, 0, D("0.00"), 0, D("3.11"), D("1.56")),
                          ("C", "Az veri", 1, D(50), D(5000), D(1000), D("20.00"), D("10.00"), D("4.00"), D(-100), D(100), 1, 1, 0, D("100.00"), 1, D("4.50"), D("0.50"))])
        gn = r["genel"]
        self.assertEqual((gn["gun"], gn["satis"], gn["iade"], gn["iade_oran"], gn["iptal"], gn["iptal_oran"], gn["fark"], gn["farkli"]), (9, D(85000), D(4600), D("5.41"), D(900), D("1.06"), D(-135), 4))
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]), [("Orta", "Kasiyer diğerlerinden ayrışıyor", "B"), ("Yüksek", "Kasiyer diğerlerinden ayrışıyor", "A")])
        self.assertEqual([(x["kasiyer"], x["tarih"].day, x["farkli"]) for x in r["gunler"][:4]], [("a", 1, True), ("A", 2, True), ("A", 3, True), ("A", 4, False)])

    def test_esikler(self):
        d = lambda **kw: {k["kasiyer"]: k["durum"] for k in main.analiz(veri(), **kw)["kasiyerler"]}   # noqa: E731
        self.assertEqual((d(seri=4, fark_gun=D(75)), d(tolerans=D(6)), d(kat=D(3)), d(asgari=50)["C"], d(asgari=401)),
                         ({"A": "Olağan", "B": "İade oranı yüksek", "C": "Az veri"}, {"A": "Olağan", "B": "İade oranı yüksek", "C": "Az veri"}, {"A": "Sürekli kasa eksiği; Sık kasa farkı", "B": "Olağan", "C": "Az veri"},
                          "İade oranı yüksek; İptal oranı yüksek; Sık kasa farkı", dict.fromkeys("ABC", "Az veri")))
        r = main.analiz([g("A", 1, -10), g("A", 2, 0), g("A", 3, -10), g("A", 5, -10), g("A", 4, 10)])
        k = r["kasiyerler"][0]
        self.assertEqual((k["seri"], k["eksik"], k["fazla"], k["farkli_oran"], k["durum"], k["diger_iade"]), (1, 3, 1, D("80.00"), "Sık kasa farkı", None))
        with self.assertRaises(ValueError):
            main.analiz([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "kasiyer_gunluk.csv", cikti)
            self.assertEqual([(k["kasiyer"], k["durum"], k["satis"], k["iade_oran"], k["iptal_oran"], k["fark"], k["farkli"], k["eksik"], k["seri"], k["diger_iade"], k["diger_iptal"]) for k in r["kasiyerler"]],
                             [("B. Demir", "İade oranı yüksek; İptal oranı yüksek; Sürekli kasa eksiği; Sık kasa farkı", D(300000), D("7.00"), D("5.20"), D(-135), 5, 5, 4, D("1.43"), D("0.56")),
                              ("A. Yılmaz", "Olağan", D(360000), D("1.50"), D("0.67"), D(2), 0, 0, 0, D("4.55"), D("3.13")), ("C. Kaya", "Olağan", D(228000), D("1.32"), D("0.39"), D(2), 2, 1, 1, D("4.00"), D("2.73"))])
            gn = r["genel"]
            self.assertEqual((gn["gun"], gn["fis"], gn["satis"], gn["iade"], gn["iade_oran"], gn["iptal"], gn["iptal_oran"], gn["fark"]), (18, D(4560), D(888000), D(29400), D("3.31"), D(18900), D("2.13"), D(-131)))
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]], [("Yüksek", "Kasiyer diğerlerinden ayrışıyor", "B. Demir")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Kasiyerler", "Günlük Kayıtlar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "k.csv"
            y.write_text("Gün;Personel;Fiş Sayısı;Ciro;İade;İptal;Fark\n01.10.2026;Ali;1.200;12.500,50;250;;-12,5\n02.10.2026;Ali;;100;;;\nx;Ali;;100;;;\n03.10.2026;Ali;;-5;;;\n03.10.2026;Ali;;5;;;y\n03.10.2026;;;5;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["tarih"], x["kasiyer"], x["fis"], x["satis"], x["iade"], x["iptal"], x["iade_adet"], x["fark"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([(date(2026, 10, 1), "Ali", D(1200), D("12500.50"), D(250), D(0), D(0), D("-12.5")), (date(2026, 10, 2), "Ali", D(0), D(100), D(0), D(0), D(0), D(0))],
                              [("Satır okunamadı", "Ali · x"), ("Satır okunamadı", "Ali · 03.10.2026"), ("Satır okunamadı", "Ali · 03.10.2026")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--kayitlar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tolerans", "10", "--kat", "1,5", "--seri", "4", "--fark-gun", "40", "--min", "500"]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tolerans", "-1"], ["--kat", "0,5"], ["--seri", "1"], ["--fark-gun", "101"], ["--min", "0"])], [2] * 5)
            self.assertEqual(main.main(["--kayitlar", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Tarih;Kasiyer\n01.10.2026;A\n", "Tarih;Kasiyer;Satış Tutarı\n01.10.2026;;5\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
