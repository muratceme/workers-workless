"""
IBNR Rezerv Tahmini (Zincirleme Merdiven) — Workers / Workless kod bloğu
Sigortacılık › Aktüerya › Aktüerya Uzmanı

Hasar gelişim üçgeninden zincirleme merdiven (chain ladder) yöntemiyle nihai hasarı ve gerçekleşmiş ancak
rapor edilmemiş / henüz ödenmemiş hasar karşılığını hesaplar:
  - Üçgen doğrudan (kümülatif veya artımlı matris) ya da hasar ödeme/ihbar dökümünden (kaza tarihi, işlem tarihi,
    tutar) yıllık veya çeyreklik olarak kurulur.
  - Gelişim faktörleri: hacim ağırlıklı (varsayılan), basit ortalama veya son N takvim yılı; isteğe bağlı kuyruk
    faktörü ve elle faktör seçimi.
  - Mack (1993) dağılımdan bağımsız standart hata: kaza yılı bazında ve toplamda.
  - Tanı tabloları: bağlantı oranları (link ratio) ve ortalamadan sapmalar, gelişmişlik yüzdesi.
İnternete bağlanmaz.

Kullanım:
    python main.py                                             # Taylor-Ashe (1983) örnek üçgeniyle
    python main.py --ucgen ucgen.xlsx                          # kümülatif üçgen (satır: kaza yılı, sütun: gelişim)
    python main.py --ucgen ucgen.xlsx --artimli --son-n 5 --kuyruk 1.02
    python main.py --hasarlar odemeler.xlsx --donem ceyrek     # dökümden üçgen kur
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def katla(s) -> str:
    s = kucuk(s)
    for a, b in zip("ıöüşğç", "iousgc"):
        s = s.replace(a, b)
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def sayi(x) -> float | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip().replace(" ", "").replace("TL", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1:
        s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        return None


def tl(x: float, basamak: int = 0) -> str:
    return f"{x:,.{basamak}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    for b in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%Y %H:%M"):
        try:
            return datetime.strptime(str(x or "").strip(), b).date()
        except ValueError:
            continue
    return None


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


# ----------------------------------------------------------------------------
# Üçgen
# ----------------------------------------------------------------------------

def ucgen_oku(yol: Path, artimli: bool) -> tuple[list[str], list[list[float]]]:
    """İlk sütun kaza dönemi, sonraki sütunlar gelişim dönemleri. Boş hücre = henüz gözlenmedi."""
    s = tablo_oku(yol)
    # Başlık ve unvan satırlarını at: veri satırının ilk hücresi dönem etiketidir (2016, 2016-Ç1 …)
    veri = [r for r in s if sum(1 for c in r if c not in (None, "")) >= 2]
    ilk = next((i for i, r in enumerate(veri) if re.match(r"^\s*\d{4}", str(r[0] or ""))), None)
    veri = veri[ilk:] if ilk is not None else veri[1:]
    etiketler, ucgen = [], []
    for r in veri:
        if r[0] in (None, ""):
            continue
        degerler = [sayi(c) for c in r[1:]]
        while degerler and degerler[-1] is None:
            degerler.pop()
        if not degerler:
            continue
        if any(v is None for v in degerler):
            raise SystemExit(f"{yol.name}: '{r[0]}' satırında aradan boş hücre var; üçgen soldan sağa kesintisiz olmalı.")
        if artimli:
            toplam, kum = 0.0, []
            for v in degerler:
                toplam += v
                kum.append(toplam)
            degerler = kum
        etiketler.append(str(r[0]).strip())
        ucgen.append(degerler)
    return etiketler, ucgen


def dokumden_ucgen(yol: Path, donem: str, degerleme: date | None) -> tuple[list[str], list[list[float]]]:
    """Hasar ödeme/ihbar dökümünden kümülatif üçgen: kaza dönemi × gelişim dönemi (işlem tarihi − kaza tarihi)."""
    s = tablo_oku(yol)
    b = [katla(x) for x in s[0]]
    bul = lambda *a: next((b.index(katla(x)) for x in a if katla(x) in b), None)  # noqa: E731
    i_kaza = bul("kaza tarihi", "hasar tarihi", "olay tarihi", "kaza")
    i_islem = bul("ödeme tarihi", "işlem tarihi", "ihbar tarihi", "rapor tarihi", "tarih")
    i_tutar = bul("tutar", "ödeme tutarı", "ödenen", "muallak değişimi", "hasar tutarı")
    if None in (i_kaza, i_islem, i_tutar):
        raise SystemExit(f"{yol.name}: Kaza Tarihi, Ödeme (İşlem) Tarihi ve Tutar sütunları gerekli. Başlıklar: {s[0]}")
    adim = 3 if donem == "ceyrek" else 12
    hucre: dict[tuple[int, int], float] = defaultdict(float)
    ilk_kaza = None
    son = None
    for r in s[1:]:
        k, t, v = tarih(r[i_kaza] if i_kaza < len(r) else None), tarih(r[i_islem] if i_islem < len(r) else None), sayi(r[i_tutar])
        if not k or not t or v is None:
            continue
        if degerleme and t > degerleme:
            continue
        kd = (k.year * 12 + k.month - 1) // adim
        td = (t.year * 12 + t.month - 1) // adim
        hucre[(kd, max(0, td - kd))] += v
        ilk_kaza = kd if ilk_kaza is None else min(ilk_kaza, kd)
        son = td if son is None else max(son, td)
    if not hucre:
        raise SystemExit(f"{yol.name}: geçerli satır bulunamadı.")
    son_kaza = max(kd for kd, _ in hucre)
    etiketler, ucgen = [], []
    for kd in range(ilk_kaza, son_kaza + 1):
        n = son - kd + 1
        kum, toplam = [], 0.0
        for g in range(n):
            toplam += hucre.get((kd, g), 0.0)
            kum.append(toplam)
        yil, ay0 = divmod(kd * adim, 12)
        etiketler.append(str(yil) if adim == 12 else f"{yil}-Ç{ay0 // 3 + 1}")
        ucgen.append(kum)
    return etiketler, ucgen


# ----------------------------------------------------------------------------
# Zincirleme merdiven ve Mack
# ----------------------------------------------------------------------------

def bag_oranlari(C: list[list[float]]) -> list[list[float | None]]:
    return [[(row[k + 1] / row[k]) if row[k] else None for k in range(len(row) - 1)] for row in C]


def faktorler(C: list[list[float]], yontem: str = "hacim", son_n: int | None = None) -> list[float]:
    n = max(len(r) for r in C)
    f = []
    for k in range(n - 1):
        satirlar = [i for i in range(len(C)) if len(C[i]) > k + 1]
        if son_n:
            satirlar = satirlar[-son_n:]                     # en son N kaza yılı = son N takvim köşegeni
        if yontem == "basit":
            oranlar = [C[i][k + 1] / C[i][k] for i in satirlar if C[i][k]]
            f.append(sum(oranlar) / len(oranlar) if oranlar else 1.0)
        else:
            pay, payda = sum(C[i][k + 1] for i in satirlar), sum(C[i][k] for i in satirlar)
            f.append(pay / payda if payda else 1.0)
    return f


def mack_sigma(C: list[list[float]], f: list[float]) -> list[float]:
    """σ²_k = 1/(n−1) Σ C_ik (C_i,k+1/C_ik − f_k)²; son faktör için Mack'in dışdeğerlemesi."""
    n = max(len(r) for r in C)
    s2: list[float] = []
    for k in range(n - 1):
        satirlar = [i for i in range(len(C)) if len(C[i]) > k + 1 and C[i][k]]
        if len(satirlar) > 1:
            s2.append(sum(C[i][k] * (C[i][k + 1] / C[i][k] - f[k]) ** 2 for i in satirlar) / (len(satirlar) - 1))
        elif len(s2) >= 2 and s2[-2] > 0:
            s2.append(min(s2[-1] ** 2 / s2[-2], s2[-2], s2[-1]))
        else:
            s2.append(0.0)
    return s2


def hesapla(etiketler: list[str], C: list[list[float]], yontem: str = "hacim", son_n: int | None = None,
            kuyruk: float = 1.0, secili: dict[int, float] | None = None) -> dict:
    n = max(len(r) for r in C)
    if len(C) < 2 or n < 2:
        raise SystemExit("Üçgen en az 2 kaza dönemi ve 2 gelişim dönemi içermeli.")
    f_hesap = faktorler(C, yontem, son_n)
    f = list(f_hesap)
    for k, v in (secili or {}).items():
        if 0 <= k < len(f):
            f[k] = v
    s2 = mack_sigma(C, f)
    tam = [list(r) for r in C]
    for i in range(len(C)):
        for k in range(len(C[i]) - 1, n - 1):
            tam[i].append(tam[i][k] * f[k])
    cdf = [1.0] * n
    for k in range(n - 2, -1, -1):
        cdf[k] = cdf[k + 1] * f[k]
    cdf = [c * kuyruk for c in cdf]
    son_gozlem = [r[-1] for r in C]
    nihai = [tam[i][-1] * kuyruk for i in range(len(C))]
    rezerv = [nihai[i] - son_gozlem[i] for i in range(len(C))]

    def S(k):
        return sum(C[j][k] for j in range(len(C)) if len(C[j]) > k + 1)

    mse = []
    for i in range(len(C)):
        a = len(C[i]) - 1
        mse.append(tam[i][-1] ** 2 * sum(s2[k] / f[k] ** 2 * (1 / tam[i][k] + 1 / S(k)) for k in range(a, n - 1)
                                         if tam[i][k] and S(k)))
    toplam_mse = sum(mse)
    for i in range(len(C)):
        a = len(C[i]) - 1
        sonraki = sum(tam[j][-1] for j in range(i + 1, len(C)))
        toplam_mse += tam[i][-1] * sonraki * sum(2 * s2[k] / f[k] ** 2 / S(k) for k in range(a, n - 1) if S(k))
    se = [math.sqrt(m) for m in mse]
    uyarilar = []
    if kuyruk != 1.0:
        uyarilar.append("Kuyruk faktörü kullanıldı: Mack standart hatası kuyruk belirsizliğini içermez.")
    if secili:
        uyarilar.append(f"Elle seçilen faktörler: {', '.join(f'{k + 1}→{k + 2}: {v}' for k, v in secili.items())}")
    return {"etiketler": etiketler, "C": C, "tam": tam, "f": f, "f_hesap": f_hesap, "s2": s2, "cdf": cdf,
            "son": son_gozlem, "nihai": nihai, "rezerv": rezerv, "se": se, "toplam_rezerv": sum(rezerv),
            "toplam_se": math.sqrt(toplam_mse), "oranlar": bag_oranlari(C), "yontem": yontem, "son_n": son_n,
            "kuyruk": kuyruk, "uyarilar": uyarilar}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
TAHMIN = PatternFill("solid", fgColor="EDE7FF")
SAPMA = PatternFill("solid", fgColor="FDE2E1")
PARA = "#,##0"
ORAN = "0.0000"


def _baslik(ws, satir=1):
    for c in ws[satir]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI


def rapor(s: dict, cikti: Path, tur: str):
    n = max(len(r) for r in s["C"])
    wb = Workbook()
    o = wb.active
    o.title = "Özet"
    o.append([f"IBNR rezerv tahmini — zincirleme merdiven ({tur})"])
    o["A1"].font = Font(bold=True, size=12)
    o.append(["Kaza Dönemi", "Son Gözlenen (kümülatif)", "Gelişmişlik %", "CDF (nihaiye)", "Nihai Hasar", "IBNR / Rezerv",
              "Mack Std. Hata", "Değişim Katsayısı"])
    _baslik(o, 2)
    for i, e in enumerate(s["etiketler"]):
        k = len(s["C"][i]) - 1
        o.append([e, s["son"][i], 1 / s["cdf"][k], s["cdf"][k], s["nihai"][i], s["rezerv"][i], s["se"][i],
                  (s["se"][i] / s["rezerv"][i]) if s["rezerv"][i] else None])
        r = o.max_row
        for c in (2, 5, 6, 7):
            o.cell(r, c).number_format = PARA
        o.cell(r, 3).number_format = "0.0%"
        o.cell(r, 4).number_format = ORAN
        o.cell(r, 8).number_format = "0.0%"
    o.append(["TOPLAM", sum(s["son"]), None, None, sum(s["nihai"]), s["toplam_rezerv"], s["toplam_se"],
              s["toplam_se"] / s["toplam_rezerv"] if s["toplam_rezerv"] else None])
    r = o.max_row
    for c in range(1, 9):
        o.cell(r, c).font = Font(bold=True)
    for c in (2, 5, 6, 7):
        o.cell(r, c).number_format = PARA
    o.cell(r, 8).number_format = "0.0%"
    o.append([])
    yontem = {"hacim": "hacim ağırlıklı ortalama", "basit": "basit ortalama"}[s["yontem"]]
    o.append(["Faktör yöntemi", yontem + (f", son {s['son_n']} kaza dönemi" if s["son_n"] else ", tüm dönemler")])
    o.append(["Kuyruk faktörü", s["kuyruk"]])
    for u in s["uyarilar"]:
        o.append(["Uyarı", u])
    o.append(["Not", "Mack standart hatası, kaza dönemi rezervinin belirsizliğini ölçer; toplam hata dönemler arası "
                     "korelasyonu içerir. Büyük hasarlar, katastrofik olaylar ve enflasyon değişimleri üçgeni bozabilir; "
                     "bu durumlarda ayıklama veya ayrı analiz gerekir."])
    for j, w in enumerate((14, 20, 13, 13, 18, 18, 16, 14), 1):
        o.column_dimensions[get_column_letter(j)].width = w
    g = BarChart()
    g.title, g.height, g.width = "Kaza dönemine göre IBNR", 8, 20
    g.add_data(Reference(o, min_col=6, min_row=2, max_row=2 + len(s["etiketler"])), titles_from_data=True)
    g.set_categories(Reference(o, min_col=1, min_row=3, max_row=2 + len(s["etiketler"])))
    o.add_chart(g, "J2")

    f = wb.create_sheet("Faktörler")
    f.append(["Gelişim", "Hesaplanan Faktör", "Kullanılan Faktör", "CDF", "σ² (Mack)"])
    _baslik(f)
    for k in range(n - 1):
        f.append([f"{k + 1}→{k + 2}", s["f_hesap"][k], s["f"][k], s["cdf"][k], s["s2"][k]])
        for c in (2, 3, 4):
            f.cell(f.max_row, c).number_format = ORAN
        f.cell(f.max_row, 5).number_format = "#,##0.00"
        if s["f"][k] != s["f_hesap"][k]:
            f.cell(f.max_row, 3).fill = SAPMA
    f.append([f"{n}→nihai (kuyruk)", None, s["kuyruk"], s["cdf"][n - 1], None])
    for j, w in enumerate((16, 18, 18, 12, 18), 1):
        f.column_dimensions[get_column_letter(j)].width = w

    u = wb.create_sheet("Üçgen")
    u.append(["Kaza Dönemi"] + [f"Gelişim {k + 1}" for k in range(n)] + ["Nihai"])
    _baslik(u)
    for i, e in enumerate(s["etiketler"]):
        u.append([e] + s["tam"][i] + [s["nihai"][i]])
        for k in range(n + 1):
            u.cell(u.max_row, k + 2).number_format = PARA
            if k >= len(s["C"][i]):
                u.cell(u.max_row, k + 2).fill = TAHMIN
    u.append([])
    u.append(["Mor hücreler zincirleme merdivenle tahmin edilen değerlerdir."])
    u.column_dimensions["A"].width = 13
    for k in range(n + 1):
        u.column_dimensions[get_column_letter(k + 2)].width = 14
    u.freeze_panes = "B2"

    b = wb.create_sheet("Bağlantı Oranları")
    b.append(["Kaza Dönemi"] + [f"{k + 1}→{k + 2}" for k in range(n - 1)])
    _baslik(b)
    for i, e in enumerate(s["etiketler"]):
        b.append([e] + s["oranlar"][i])
        for k, v in enumerate(s["oranlar"][i]):
            c = b.cell(b.max_row, k + 2)
            c.number_format = ORAN
            sd = math.sqrt(s["s2"][k] / s["C"][i][k]) if s["C"][i][k] and s["s2"][k] > 0 else None
            if v is not None and sd and abs(v - s["f"][k]) > 2 * sd:
                c.fill = SAPMA                       # Mack modeline göre 2 standart sapmadan uzak
    b.append(["Kullanılan"] + s["f"])
    for k in range(n - 1):
        b.cell(b.max_row, k + 2).number_format = ORAN
        b.cell(b.max_row, k + 2).font = Font(bold=True)
    b.append([])
    b.append(["Kırmızı: kullanılan faktörden 2 standart sapmadan fazla uzak bağlantı oranı (√(σ²/C) ölçeğinde); "
              "büyük hasar veya süreç değişikliği olabilir."])
    for k in range(n):
        b.column_dimensions[get_column_letter(k + 1)].width = 13
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def secili_coz(degerler: list[str] | None) -> dict[int, float]:
    sonuc = {}
    for d in degerler or []:
        m = re.match(r"^(\d+)=([\d.,]+)$", d.strip())
        if not m:
            raise SystemExit(f"--faktor biçimi gelişim=değer olmalı (ör. 7=1,05): {d}")
        sonuc[int(m.group(1)) - 1] = sayi(m.group(2))
    return sonuc


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description="Hasar gelişim üçgeninden zincirleme merdiven yöntemiyle IBNR ve Mack standart hatasını hesaplar.")
    kaynak = ap.add_mutually_exclusive_group()
    kaynak.add_argument("--ucgen", type=Path, help="Üçgen (.xlsx/.csv): ilk sütun kaza dönemi, sonraki sütunlar gelişim dönemleri")
    kaynak.add_argument("--hasarlar", type=Path, help="Döküm (.xlsx/.csv): Kaza Tarihi, Ödeme/İşlem Tarihi, Tutar")
    ap.add_argument("--artimli", action="store_true", help="Üçgen artımlı (dönem içi) tutarlarla verildi")
    ap.add_argument("--donem", choices=["yil", "ceyrek"], default="yil", help="Dökümden üçgen kurarken dönem (varsayılan yıl)")
    ap.add_argument("--degerleme", help="Dökümde değerleme tarihi (GG.AA.YYYY); sonrası dikkate alınmaz")
    ap.add_argument("--yontem", choices=["hacim", "basit"], default="hacim", help="Faktör ortalaması (varsayılan hacim ağırlıklı)")
    ap.add_argument("--son-n", type=int, help="Faktörleri yalnız son N kaza döneminden hesapla")
    ap.add_argument("--kuyruk", type=float, default=1.0, help="Kuyruk faktörü (varsayılan 1)")
    ap.add_argument("--faktor", nargs="*", help="Elle faktör seçimi, ör. --faktor 8=1,07 9=1,015 (gelişim=değer)")
    ap.add_argument("--tur", default="ödenen hasar", help="Rapor başlığında üçgen türü (ödenen / gerçekleşen hasar)")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "ibnr_zincirleme_merdiven.xlsx")
    a = ap.parse_args(argv)
    if a.hasarlar:
        etiketler, C = dokumden_ucgen(a.hasarlar, a.donem, tarih(a.degerleme) if a.degerleme else None)
    else:
        etiketler, C = ucgen_oku(a.ucgen or BURASI / "ornek_veri" / "taylor_ashe_kumulatif.csv", a.artimli)
    s = hesapla(etiketler, C, a.yontem, a.son_n, a.kuyruk, secili_coz(a.faktor))
    rapor(s, a.cikti, a.tur)
    print(f"[OK] {len(C)} kaza dönemi · toplam IBNR {tl(s['toplam_rezerv'])} · Mack std. hata {tl(s['toplam_se'])} "
          f"(değişim katsayısı %{tl(s['toplam_se'] / s['toplam_rezerv'] * 100, 1)})")
    print("    faktörler: " + " ".join(f"{x:.4f}" for x in s["f"]))
    for u in s["uyarilar"]:
        print(f"[!] {u}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
