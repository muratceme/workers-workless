"""
Güvenlik Raporlama Sistemi Olay Analizi — Gerisi Bende kod bloğu
Kalite Yönetimi › Kalite Birim Sorumlusu
Hasta ve çalışan güvenliği olay bildirimlerini tür, kategori, birim, zarar düzeyi ve önlenebilirlik açısından analiz eder.
  Dağılımlar        = olay türü, kategori, birim, zarar düzeyi ve ay bazında bildirim sayısı ve payı
  Ramak kala oranı  = ramak kala bildirimi ÷ toplam bildirim × 100   (yüksek olması bildirim kültürünün işaretidir)
  Önlenebilir oranı = önlenebilir ÷ önlenebilirliği değerlendirilmiş bildirim × 100
  1000 hasta gününe = birim bildirimi ÷ birim hasta günü × 1000   (--hasta-gunu dosyası verilirse)
  Uyarı             = ciddi zarar ya da ölümle sonuçlanan olay · bu olaylarda DÖF numarası yok · --gun (30) günden uzun süredir açık bildirim
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"no": ("bildirim no", "olay no", "kayit no", "no"), "tarih": ("olay tarihi", "tarih", "bildirim tarihi"), "birim": ("birim", "bolum", "servis", "olay yeri"),
         "tur": ("olay turu", "bildirim turu", "guvenlik turu", "tur"), "kategori": ("olay kategorisi", "kategori", "olay tipi", "alt tur"), "zarar": ("zarar duzeyi", "zarar", "sonuc", "siddet"),
         "onlenebilir": ("onlenebilir", "onlenebilirlik", "onlenebilir mi"), "kok": ("kok neden", "kok neden kategorisi", "neden"), "dof": ("dof no", "dof", "duzeltici faaliyet no"),
         "kapanis": ("kapanis tarihi", "kapatma tarihi", "kapanis")}
HG = {"birim": ("birim", "bolum", "servis"), "hg": ("hasta gunu", "hasta gun sayisi", "yatak gunu")}
ZARAR = (("Ramak kala", ("ramak", "near miss", "olmadan onlendi")), ("Zarar yok", ("zarar yok", "zararsiz", "etkisiz")), ("Hafif", ("hafif", "minor", "dusuk")),
         ("Orta", ("orta", "moderate")), ("Ciddi", ("ciddi", "agir", "kalici", "major", "yuksek")), ("Ölüm", ("olum", "vefat", "exitus")))
CIDDI = ("Ciddi", "Ölüm")
SIFIR = Decimal(0)


def _zarar(x) -> str | None:
    k = katla(x)
    if not k:
        return "Belirtilmemiş"
    return next((ad for ad, anahtar in ZARAR if any(a in k for a in anahtar)), None)


def _tur(x) -> str:
    k = katla(x)
    if "hasta" in k:
        return "Hasta güvenliği"
    if "calisan" in k or "personel" in k:
        return "Çalışan güvenliği"
    return metin(x) or "Belirtilmemiş"


def _evet(x) -> str:
    k = katla(x)
    return "Evet" if k in ("evet", "e", "x", "1", "onlenebilir") else "Hayır" if k in ("hayir", "h", "0", "onlenemez") else "Değerlendirilmedi"


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("tarih", "birim")):
        no, birim = metin(r["no"]), metin(r["birim"])
        if not no and not birim:
            continue
        t, z = tarih(r["tarih"]), _zarar(r["zarar"])
        kap = tarih(r["kapanis"]) if metin(r["kapanis"]) else None
        kim = no or f"Satır {r['_satir']}"
        if t is None or z is None or (metin(r["kapanis"]) and kap is None):
            uy.append(uyari("Yüksek", "Satır okunamadı", kim, f"Satır {r['_satir']}: tarih, kapanış tarihi ya da zarar düzeyi ('{metin(r['zarar'])}') okunamadı. Bildirim alınmadı."))
            continue
        if no and katla(no) in gorulen:
            uy.append(uyari("Orta", "Bildirim iki kez yazılmış", kim, f"Satır {r['_satir']}: ikinci satır alınmadı."))
            continue
        if no:
            gorulen.add(katla(no))
        l.append({"no": no, "tarih": t, "birim": birim or "Belirtilmemiş", "tur": _tur(r["tur"]), "kategori": metin(r["kategori"]) or "Belirtilmemiş", "zarar": z,
                  "onlenebilir": _evet(r["onlenebilir"]), "kok": metin(r["kok"]), "dof": metin(r["dof"]), "kapanis": kap, "ay": f"{t:%Y-%m}"})
    return l, uy


def oku_hg(yol: Path) -> tuple[dict, list[dict]]:
    uy, d = [], {}
    for r in kayitlar(yol, HG, ("birim", "hg")):
        b = metin(r["birim"])
        if not b:
            continue
        v = para(r["hg"])
        if v is None or v <= 0:
            uy.append(uyari("Yüksek", "Hasta günü okunamadı", b, f"Satır {r['_satir']}: hasta günü pozitif sayı olmalı."))
            continue
        d[katla(b)] = v
    return d, uy


def _dagilim(l: list[dict], alan: str, sira: tuple | None = None) -> list[dict]:
    g = {}
    for o in l:
        g.setdefault(katla(o[alan]), []).append(o)
    n = len(l)
    satir = [{"ad": x[0][alan], "adet": len(x), "pay": yuvarla(Decimal(len(x)) * 100 / n), "ramak": sum(1 for o in x if o["zarar"] == "Ramak kala"),
              "ciddi": sum(1 for o in x if o["zarar"] in CIDDI), "onlenebilir": sum(1 for o in x if o["onlenebilir"] == "Evet"), "acik": sum(1 for o in x if o["kapanis"] is None)} for x in g.values()]
    if sira:
        return sorted(satir, key=lambda s: sira.index(s["ad"]) if s["ad"] in sira else len(sira))
    return sorted(satir, key=lambda s: (s["ad"] if alan == "ay" else "", -s["adet"], katla(s["ad"])))


def analiz(olaylar: list[dict], gun: date, acik_gun: int = 30, hasta_gunu: dict | None = None) -> dict:
    if not olaylar:
        raise ValueError("Olay bildirimi yok.")
    uy = []
    n = len(olaylar)
    degerlendirilen = [o for o in olaylar if o["onlenebilir"] != "Değerlendirilmedi"]
    ozet = {"bildirim": n, "ramak": sum(1 for o in olaylar if o["zarar"] == "Ramak kala"), "ciddi": sum(1 for o in olaylar if o["zarar"] in CIDDI),
            "acik": sum(1 for o in olaylar if o["kapanis"] is None), "degerlendirilen": len(degerlendirilen), "onlenebilir": sum(1 for o in degerlendirilen if o["onlenebilir"] == "Evet")}
    ozet["ramak_oran"] = yuvarla(Decimal(ozet["ramak"]) * 100 / n)
    ozet["onlenebilir_oran"] = yuvarla(Decimal(ozet["onlenebilir"]) * 100 / len(degerlendirilen)) if degerlendirilen else None
    birim = _dagilim(olaylar, "birim")
    for b in birim:
        hg = (hasta_gunu or {}).get(katla(b["ad"]))
        b["hg"], b["oran"] = hg, yuvarla(Decimal(b["adet"]) * 1000 / hg) if hg else None
    for o in sorted(olaylar, key=lambda o: o["tarih"]):
        if o["zarar"] in CIDDI:
            uy.append(uyari("Yüksek", "Ciddi sonuçlu olay", o["no"] or o["birim"], f"{o['tarih']:%d.%m.%Y} · {o['birim']} · {o['kategori']} · {o['zarar']}. Kök neden analizi ve bildirim yükümlülükleri kontrol edilmeli."))
            if not o["dof"]:
                uy.append(uyari("Yüksek", "Ciddi olayda DÖF yok", o["no"] or o["birim"], "Düzeltici / önleyici faaliyet numarası yazılmamış."))
        if o["kapanis"] is None and (gun - o["tarih"]).days > acik_gun:
            uy.append(uyari("Orta", "Uzun süredir açık bildirim", o["no"] or o["birim"], f"{o['tarih']:%d.%m.%Y} tarihli bildirim {(gun - o['tarih']).days} gündür kapanmadı (eşik {acik_gun})."))
        if o["kapanis"] and o["kapanis"] < o["tarih"]:
            uy.append(uyari("Orta", "Kapanış olaydan önce", o["no"] or o["birim"], f"Olay {o['tarih']:%d.%m.%Y}, kapanış {o['kapanis']:%d.%m.%Y}."))
    if ozet["degerlendirilen"] < n:
        uy.append(uyari("Bilgi", "Önlenebilirlik değerlendirilmemiş", f"{n - ozet['degerlendirilen']} bildirim", "Önlenebilir sütunu boş ya da tanınmayan bildirimler oranın dışında tutuldu."))
    acik = sorted(({**o, "gun": (gun - o["tarih"]).days} for o in olaylar if o["kapanis"] is None), key=lambda o: (-o["gun"], o["no"]))
    return {"ozet": ozet, "tur": _dagilim(olaylar, "tur"), "kategori": _dagilim(olaylar, "kategori"), "birim": birim, "zarar": _dagilim(olaylar, "zarar", tuple(a for a, _ in ZARAR) + ("Belirtilmemiş",)),
            "ay": _dagilim(olaylar, "ay"), "kok": _dagilim([o for o in olaylar if o["kok"]], "kok") if any(o["kok"] for o in olaylar) else [], "acik": acik, "olaylar": olaylar, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, gun: date) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    o = r["ozet"]
    _baslik(ws, ["Gösterge", "Değer"], (44, 14))
    for ad, v in (("Bildirim sayısı", o["bildirim"]), ("Ramak kala", o["ramak"]), ("Ramak kala oranı %", o["ramak_oran"]), ("Ciddi zarar / ölüm", o["ciddi"]),
                  ("Önlenebilirliği değerlendirilen", o["degerlendirilen"]), ("Önlenebilir", o["onlenebilir"]), ("Önlenebilir oranı %", o["onlenebilir_oran"]), ("Açık bildirim", o["acik"]),
                  ("Rapor tarihi", gun)):
        _satir(ws, [ad, v], para_s=(2,) if "%" in ad else (), tarih_s=(2,) if isinstance(v, date) else ())
    ws = wb.create_sheet("Dağılımlar")
    _baslik(ws, ["Kırılım", "Ad", "Bildirim", "Pay %", "Ramak Kala", "Ciddi / Ölüm", "Önlenebilir", "Açık", "Hasta Günü", "1000 Hasta Gününe"], (12, 30, 9, 8, 9, 9, 10, 7, 11, 11), "C2")
    for ad, satirlar in (("Olay türü", r["tur"]), ("Zarar düzeyi", r["zarar"]), ("Kategori", r["kategori"]), ("Birim", r["birim"]), ("Ay", r["ay"]), ("Kök neden", r["kok"])):
        for s in satirlar:
            _satir(ws, [ad, s["ad"], s["adet"], s["pay"], s["ramak"], s["ciddi"], s["onlenebilir"], s["acik"], s.get("hg"), s.get("oran")], para_s=(4, 10), dolgu={6: KIRMIZI if s["ciddi"] else None})
    ws = wb.create_sheet("Açık Bildirimler")
    _baslik(ws, ["Bildirim No", "Olay Tarihi", "Geçen Gün", "Birim", "Tür", "Kategori", "Zarar", "DÖF No"], (12, 11, 9, 20, 18, 24, 12, 12))
    for a in r["acik"]:
        _satir(ws, [a["no"], a["tarih"], a["gun"], a["birim"], a["tur"], a["kategori"], a["zarar"], a["dof"]], tarih_s=(2,), dolgu={7: KIRMIZI if a["zarar"] in CIDDI else None})
    ws = wb.create_sheet("Bildirimler")
    _baslik(ws, ["Bildirim No", "Tarih", "Birim", "Tür", "Kategori", "Zarar", "Önlenebilir", "Kök Neden", "DÖF No", "Kapanış"], (12, 11, 20, 18, 24, 12, 11, 22, 12, 11), "B2")
    for x in sorted(r["olaylar"], key=lambda x: (x["tarih"], x["no"])):
        _satir(ws, [x["no"], x["tarih"], x["birim"], x["tur"], x["kategori"], x["zarar"], x["onlenebilir"], x["kok"], x["dof"], x["kapanis"]], tarih_s=(2, 10))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, gun: date, acik_gun: int = 30, hg_yol: Path | None = None) -> dict:
    olaylar, uy = oku(yol)
    hg, uy2 = oku_hg(hg_yol) if hg_yol else (None, [])
    r = analiz(olaylar, gun, acik_gun, hg)
    r["uyarilar"] = sirala(uy + uy2 + r["uyarilar"])
    rapor_yaz(cikti, r, gun)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Hasta ve çalışan güvenliği olay bildirimlerini tür, birim ve önlenebilirlik açısından analiz eder.")
    p.add_argument("--bildirimler", type=Path, help="Bildirim No, Olay Tarihi, Birim, Olay Türü, Kategori, Zarar Düzeyi, Önlenebilir, Kök Neden, DÖF No, Kapanış Tarihi")
    p.add_argument("--hasta-gunu", type=Path, help="İsteğe bağlı: Birim, Hasta Günü (1000 hasta gününe düşen bildirim için)")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 30.09.2026)")
    p.add_argument("--gun", type=int, default=30, help="Bu kadar günden uzun süredir açık bildirimler uyarılır (varsayılan 30)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "grs_olay_analizi.xlsx")
    a = p.parse_args(argv)
    ornek = a.bildirimler is None
    yol = a.bildirimler or ORNEK / "olay_bildirimleri.csv"
    hg = a.hasta_gunu or (ORNEK / "birim_hasta_gunu.csv" if ornek else None)
    gun = tarih(a.tarih) if a.tarih else (date(2026, 9, 30) if ornek else date.today())
    if gun is None or a.gun < 0:
        print("[X] --tarih GG.AA.YYYY, --gun sıfır ya da üstü olmalı")
        return 2
    for y in (yol, hg):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(yol, a.cikti, gun, a.gun, hg)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {o['bildirim']} bildirim · ramak kala %{o['ramak_oran']} · ciddi / ölüm {o['ciddi']} · önlenebilir %{o['onlenebilir_oran']} · açık {o['acik']}")
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
