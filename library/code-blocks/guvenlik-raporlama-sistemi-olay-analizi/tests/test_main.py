import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
GUN = date(2026, 9, 30)


def o(no, t, birim, tur, kategori, zarar, onl, kapanis=None, dof=""):
    return {"no": no, "tarih": date(2026, *t), "birim": birim, "tur": tur, "kategori": kategori, "zarar": zarar, "onlenebilir": onl, "kok": "", "dof": dof,
            "kapanis": date(2026, *kapanis) if kapanis else None, "ay": f"2026-{t[0]:02d}"}


def veri():
    return [o("O1", (9, 1), "B1", "Hasta güvenliği", "Düşme", "Hafif", "Evet", (9, 5)), o("O2", (9, 2), "B1", "Hasta güvenliği", "İlaç hatası", "Ramak kala", "Evet"),
            o("O3", (8, 3), "B2", "Çalışan güvenliği", "Kesici-delici", "Ramak kala", "Evet"), o("O4", (9, 10), "B2", "Hasta güvenliği", "düşme", "Ciddi", "Hayır"),
            o("O5", (9, 15), "b1", "Hasta güvenliği", "İlaç hatası", "Zarar yok", "Değerlendirilmedi", (9, 16))]


class AnalizTesti(unittest.TestCase):
    def test_ozet(self):
        r = main.analiz(veri(), GUN, hasta_gunu={"b1": D(300)})
        self.assertEqual(r["ozet"], {"bildirim": 5, "ramak": 2, "ciddi": 1, "acik": 3, "degerlendirilen": 4, "onlenebilir": 3, "ramak_oran": D("40.00"), "onlenebilir_oran": D("75.00")})
        self.assertEqual([(b["ad"], b["adet"], b["pay"], b["oran"]) for b in r["birim"]], [("B1", 3, D("60.00"), D("10.00")), ("B2", 2, D("40.00"), None)])
        self.assertEqual([(z["ad"], z["adet"]) for z in r["zarar"]], [("Ramak kala", 2), ("Zarar yok", 1), ("Hafif", 1), ("Ciddi", 1)])
        self.assertEqual([(k["ad"], k["adet"], k["ciddi"]) for k in r["kategori"]], [("Düşme", 2, 1), ("İlaç hatası", 2, 0), ("Kesici-delici", 1, 0)])
        self.assertEqual([(a["ad"], a["adet"]) for a in r["ay"]], [("2026-08", 1), ("2026-09", 4)])
        self.assertEqual([(a["no"], a["gun"]) for a in r["acik"]], [("O3", 58), ("O2", 28), ("O4", 20)])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Bilgi", "Önlenebilirlik değerlendirilmemiş", "1 bildirim"), ("Orta", "Uzun süredir açık bildirim", "O3"), ("Yüksek", "Ciddi olayda DÖF yok", "O4"), ("Yüksek", "Ciddi sonuçlu olay", "O4")])

    def test_cevirme(self):
        self.assertEqual([main._zarar(x) for x in ("", "Ramak Kala", "Zarar yok", "hafif zarar", "ORTA", "Kalıcı hasar", "Ölüm", "x")],
                         ["Belirtilmemiş", "Ramak kala", "Zarar yok", "Hafif", "Orta", "Ciddi", "Ölüm", None])
        self.assertEqual([main._tur(x) for x in ("Hasta Güvenliği", "çalışan", "Tesis", "")], ["Hasta güvenliği", "Çalışan güvenliği", "Tesis", "Belirtilmemiş"])
        self.assertEqual([main._evet(x) for x in ("Evet", "hayır", "", "?")], ["Evet", "Hayır", "Değerlendirilmedi", "Değerlendirilmedi"])
        r = main.analiz(veri(), GUN, acik_gun=60)
        self.assertNotIn("Uzun süredir açık bildirim", [u["tur"] for u in r["uyarilar"]])
        with self.assertRaises(ValueError):
            main.analiz([], GUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "olay_bildirimleri.csv", cikti, GUN, 30, ORNEK / "birim_hasta_gunu.csv")
            self.assertEqual(r["ozet"], {"bildirim": 74, "ramak": 21, "ciddi": 1, "acik": 27, "degerlendirilen": 58, "onlenebilir": 42, "ramak_oran": D("28.38"), "onlenebilir_oran": D("72.41")})
            self.assertEqual([(b["ad"], b["adet"], b["oran"]) for b in r["birim"]][:2], [("Yoğun Bakım", 18, D("16.36")), ("Ameliyathane", 17, None)])
            self.assertEqual([u["tur"] for u in r["uyarilar"]][:2] + [len(r["uyarilar"])], ["Ciddi olayda DÖF yok", "Ciddi sonuçlu olay", 12])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Dağılımlar", "Açık Bildirimler", "Bildirimler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, z = Path(tmp) / "b.csv", Path(tmp) / "h.csv"
            y.write_text("Olay No;Tarih;Bölüm;Sonuç;Kapatma Tarihi\n1;01.09.2026;A;Hafif;\n1;01.09.2026;A;;\n2;x;A;;\n3;01.09.2026;A;bilinmez;\n4;01.09.2026;A;;yok\n5;01.09.2026;;;\n", encoding="utf-8")
            z.write_text("Bölüm;Hasta Günü\nA;1.200\nB;0\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["zarar"], x["birim"]) for x in l], [u["tur"] for u in uy]),
                             ([("1", "Hafif", "A"), ("5", "Belirtilmemiş", "Belirtilmemiş")], ["Bildirim iki kez yazılmış", "Satır okunamadı", "Satır okunamadı", "Satır okunamadı"]))
            d, uy = main.oku_hg(z)
            self.assertEqual((d, len(uy)), ({"a": D(1200)}, 1))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--bildirimler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--hasta-gunu", str(z), "--tarih", "30.09.2026", "--gun", "10"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--gun", "-1"])], [2] * 2)
            self.assertEqual(main.main(arg + ["--hasta-gunu", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Olay No;Sonuç\n1;x\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
