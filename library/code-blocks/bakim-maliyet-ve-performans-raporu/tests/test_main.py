import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def e(no, ekipman, tur, iscilik, malzeme=0, dis=0, durus=0, saat=0, t=date(2026, 7, 1)):
    return {"no": no, "tarih": t, "ekipman": ekipman, "tur": tur, "iscilik": D(iscilik), "malzeme": D(malzeme), "dis": D(dis), "durus": D(durus), "saat": D(saat)}


def veri():
    return [e("1", "E1", "Planlı", 100, 50, 0, 2, 4), e("2", "e1", "Plansız", 200, 300, 500, 10, 8), e("3", "E1", "Plansız", 100, 0, 0, 2, 4, date(2026, 8, 1)), e("4", "E2", "Planlı", 150, 100, 0, 0, 6, date(2026, 8, 2)),
            e("5", "E2", "Belirsiz", 50, t=None)]


def k(g):
    return (g["emir"], g["iscilik"], g["malzeme"], g["dis"], g["toplam"], g["planli"], g["plansiz"], g["planli_pay"], g["planli_emir_pay"], g["planli_saat_pay"], g["ariza"], g["durus"], g["planli_durus"], g["mttr"], g["mtbf"])


class RaporTesti(unittest.TestCase):
    def test_gostergeler(self):
        r = main.raporla(veri(), calisma=D(100))
        self.assertEqual(k(r["toplam"]), (5, D(600), D(450), D(500), D(1550), D(400), D(1100), D("25.81"), D("40.00"), D("45.45"), 2, D(12), D(2), D("6.00"), D("94.00")))
        self.assertEqual([(x["ekipman"], x["pay"]) + k(x) for x in r["ekipmanlar"]],
                         [("E1", D("80.65"), 3, D(400), D(350), D(500), D(1250), D(150), D(1100), D("12.00"), D("33.33"), D("25.00"), 2, D(12), D(2), D("6.00"), D("44.00")),
                          ("E2", D("19.35"), 2, D(200), D(100), D(0), D(300), D(250), D(0), D("83.33"), D("50.00"), D("100.00"), 0, D(0), D(0), None, None)])
        self.assertEqual([(x["tur"], x["emir"], x["toplam"]) for x in r["turler"]], [("Planlı", 2, D(400)), ("Plansız", 2, D(1100)), ("Belirsiz", 1, D(50))])
        self.assertEqual([(x["ay"], x["emir"], x["toplam"], x["ariza"]) for x in r["aylar"]], [("2026-07", 2, D(1150), 1), ("2026-08", 2, D(350), 1)])
        self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]], [("Orta", "Planlı bakım payı düşük", "Toplam")])

    def test_esikler(self):
        u = lambda **kw: sorted((x["tur"], x["kim"]) for x in main.raporla(veri(), **kw)["uyarilar"])   # noqa: E731
        self.assertEqual((u(planli_hedef=D("25.81")), u(planli_hedef=D("25.82"), ariza_esik=2)), ([], [("Planlı bakım payı düşük", "Toplam"), ("Tekrarlayan arıza", "E1")]))
        g = main.raporla(veri())["toplam"]
        self.assertEqual((g["mtbf"], main.raporla(veri(), calisma=D(6))["ekipmanlar"][0]["mtbf"], main.raporla([e("1", "E", "Planlı", 0)])["toplam"]["planli_pay"]), (None, None, None))
        with self.assertRaises(ValueError):
            main.raporla([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "bakim_is_emirleri.csv", cikti)
            self.assertEqual(k(r["toplam"]), (12, D(36500), D(82200), D(45000), D(163700), D(43000), D(120700), D("26.27"), D("50.00"), D("39.73"), 6, D(42), D("9.5"), D("7.00"), None))
            self.assertEqual([(x["ekipman"], x["toplam"], x["pay"], x["ariza"], x["durus"], x["mttr"], x["planli_pay"]) for x in r["ekipmanlar"]],
                             [("Pres 1", D(76500), D("46.73"), 3, D(22), D("7.33"), D("19.61")), ("CNC 1", D(67500), D("41.23"), 1, D(14), D("14.00"), D("27.41")), ("Kompresör", D(12500), D("7.64"), 1, D(4), D("4.00"), D("40.00")),
                              ("Pres 2", D(7200), D("4.40"), 1, D(2), D("2.00"), D("62.50"))])
            self.assertEqual([(x["ay"], x["toplam"]) for x in r["aylar"]], [("2026-07", D(74000)), ("2026-08", D(44700)), ("2026-09", D(45000))])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]], [("Orta", "Planlı bakım payı düşük", "Toplam"), ("Orta", "Tekrarlayan arıza", "Pres 1")])
            self.assertEqual(main.calistir(ORNEK / "bakim_is_emirleri.csv", cikti, calisma=D(1500))["ekipmanlar"][0]["mtbf"], D("492.67"))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Ekipmanlar", "Bakım Türleri", "Aylar", "İş Emirleri", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "b.csv"
            y.write_text("Bakım No;Bakım Tarihi;Makine;Tür;İşçilik;Yedek Parça Maliyeti;Taşeron Maliyeti;Duruş;Adam Saat\nB1;01.07.2026;M1;ÖNLEYİCİ BAKIM;1.250,50;500;;2,5;4\n;;M1;Arızi;100;;;;\nB3;;M2;Revizyon;10;;;;\nB4;;M2;Arıza;x;;;;\n"
                         "B5;;M2;Arıza;-5;;;;\nB6;;;Arıza;5;;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["tarih"], x["ekipman"], x["tur"], x["iscilik"], x["malzeme"], x["dis"], x["durus"], x["saat"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("B1", date(2026, 7, 1), "M1", "Planlı", D("1250.50"), D(500), D(0), D("2.5"), D(4)), ("Satır 3", None, "M1", "Plansız", D(100), D(0), D(0), D(0), D(0)), ("B3", None, "M2", "Belirsiz", D(10), D(0), D(0), D(0), D(0))],
                              [("Bakım türü tanınmadı", "B3"), ("İş emri okunamadı", "B4"), ("İş emri okunamadı", "B5")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--emirler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--planli", "80", "--ariza", "1", "--calisma-saat", "1.500"]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--planli", "101"], ["--ariza", "0"], ["--calisma-saat", "0"], ["--calisma-saat", "x"])], [2] * 4)
            self.assertEqual(main.main(["--emirler", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Ekipman;Tarih\nA;B\n", "Ekipman;Bakım Türü\n;Arıza\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
