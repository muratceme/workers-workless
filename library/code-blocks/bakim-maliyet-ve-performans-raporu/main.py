"""
Bakım Maliyet ve Performans Raporu — Gerisi Bende kod bloğu
Bakım Onarım › Bakım Müdürü

Bakım iş emirlerinden işçilik, malzeme ve dış hizmet maliyetlerini planlı / plansız ayrımıyla raporlar.
  Toplam maliyet          = işçilik + malzeme + dış hizmet
  Planlı bakım            = periyodik, önleyici, kestirimci, planlı;   plansız = arıza, acil, düzeltici
  Planlı bakım payı %     = planlı bakım maliyeti ÷ toplam maliyet × 100   (iş emri adedi ve işçilik saati payı ayrıca verilir)
  Arıza sayısı            = plansız iş emri adedi
  Ortalama onarım süresi  = plansız duruş saati ÷ arıza sayısı                                   (MTTR)
  Arızalar arası süre     = (--calisma-saat − plansız duruş saati) ÷ arıza sayısı                (MTBF; çalışma saati verilirse)
  Uyarı                   = planlı bakım payı --planli (%70) altında · bir ekipmanda arıza sayısı --ariza (3) ve üstünde
Maliyetleri siz verirsiniz; duruşun üretim kaybı maliyeti rapora dahil değildir.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"no": ("is emri no", "is emri", "bakim no"), "tarih": ("tarih", "is emri tarihi", "bakim tarihi"), "ekipman": ("ekipman", "makine", "ekipman adi"), "tur": ("bakim turu", "tur", "is emri turu"),
         "iscilik": ("iscilik maliyeti", "iscilik", "iscilik tutari"), "malzeme": ("malzeme maliyeti", "malzeme", "yedek parca maliyeti"), "dis": ("dis hizmet maliyeti", "dis hizmet", "taseron maliyeti"),
         "durus": ("durus suresi saat", "durus suresi", "durus saat", "durus"), "saat": ("iscilik saati", "adam saat", "calisma saati")}
TUR = {"planli": "Planlı", "periyodik": "Planlı", "onleyici": "Planlı", "kestirimci": "Planlı", "koruyucu": "Planlı", "plansiz": "Plansız", "ariza": "Plansız", "arizi": "Plansız", "acil": "Plansız", "duzeltici": "Plansız"}
MALIYET = ("iscilik", "malzeme", "dis")
SIFIR = Decimal(0)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("ekipman", "tur")):
        ekipman, no = metin(r["ekipman"]), metin(r["no"]) or f"Satır {r['_satir']}"
        if not ekipman:
            continue
        d = {a: para(r[a]) if metin(r[a]) else SIFIR for a in MALIYET + ("durus", "saat")}
        if any(v is None or v < 0 for v in d.values()):
            uy.append(uyari("Yüksek", "İş emri okunamadı", no, f"Satır {r['_satir']}: maliyet ya da süre değerlerinden biri okunamadı. Satır alınmadı."))
            continue
        tur = TUR.get(katla(r["tur"]).split(" ")[0] if katla(r["tur"]) else "")
        if tur is None:
            uy.append(uyari("Orta", "Bakım türü tanınmadı", no, f"'{metin(r['tur'])}' planlı ya da plansız olarak okunamadı; 'Belirsiz' sayıldı ve planlı bakım payına girmedi."))
        l.append({"no": no, "tarih": tarih(r["tarih"]), "ekipman": ekipman, "tur": tur or "Belirsiz", **d})
    return l, uy


def _oran(pay, payda) -> Decimal | None:
    return yuvarla(Decimal(pay) * 100 / Decimal(payda)) if payda else None


def gosterge(l: list[dict], calisma: Decimal | None = None) -> dict:
    t = {a: sum((e[a] for e in l), SIFIR) for a in MALIYET + ("saat",)}
    planli, ariza = [e for e in l if e["tur"] == "Planlı"], [e for e in l if e["tur"] == "Plansız"]
    toplam, p_tutar, durus = sum(t[a] for a in MALIYET), sum((e["toplam"] for e in planli), SIFIR), sum((e["durus"] for e in ariza), SIFIR)
    return {"emir": len(l), **t, "toplam": toplam, "planli": p_tutar, "plansiz": sum((e["toplam"] for e in ariza), SIFIR), "planli_pay": _oran(p_tutar, toplam), "planli_emir_pay": _oran(len(planli), len(l)),
            "planli_saat_pay": _oran(sum((e["saat"] for e in planli), SIFIR), t["saat"]), "ariza": len(ariza), "durus": durus, "planli_durus": sum((e["durus"] for e in planli), SIFIR),
            "mttr": yuvarla(durus / len(ariza)) if ariza else None, "mtbf": yuvarla((calisma - durus) / len(ariza)) if ariza and calisma is not None and calisma > durus else None}


def raporla(emirler: list[dict], planli_hedef: Decimal = Decimal(70), ariza_esik: int = 3, calisma: Decimal | None = None) -> dict:
    if not emirler:
        raise ValueError("Raporlanacak bakım iş emri yok.")
    for e in emirler:
        e["toplam"] = sum(e[a] for a in MALIYET)
    grup = {}
    for e in emirler:
        grup.setdefault(katla(e["ekipman"]), []).append(e)
    toplam, uy = gosterge(emirler, calisma * len(grup) if calisma is not None else None), []
    ekipmanlar = sorted(({"ekipman": g[0]["ekipman"], **gosterge(g, calisma)} for g in grup.values()), key=lambda x: (-x["toplam"], katla(x["ekipman"])))
    for x in ekipmanlar:
        x["pay"] = _oran(x["toplam"], toplam["toplam"])
        if x["ariza"] >= ariza_esik:
            uy.append(uyari("Orta", "Tekrarlayan arıza", x["ekipman"], f"Dönemde {x['ariza']} arıza, {x['durus']} saat plansız duruş, plansız bakım maliyeti {x['plansiz']:,.2f}. Kök neden analizi ve bakım planının gözden geçirilmesi önerilir."))
    if toplam["planli_pay"] is not None and toplam["planli_pay"] < planli_hedef:
        uy.append(uyari("Orta", "Planlı bakım payı düşük", "Toplam", f"Bakım maliyetinin %{toplam['planli_pay']} kadarı planlı (hedef %{planli_hedef}); iş emirlerinin %{toplam['planli_emir_pay']} kadarı planlı."))
    turler = [{"tur": t, **gosterge(g)} for t in ("Planlı", "Plansız", "Belirsiz") if (g := [e for e in emirler if e["tur"] == t])]
    ay = {}
    for e in emirler:
        if e["tarih"]:
            ay.setdefault(f"{e['tarih']:%Y-%m}", []).append(e)
    return {"emirler": sorted(emirler, key=lambda e: (e["tarih"] or date.min, e["no"])), "toplam": toplam, "ekipmanlar": ekipmanlar, "turler": turler, "aylar": [{"ay": k, **gosterge(g)} for k, g in sorted(ay.items())], "uyarilar": uy}


BASLIK = ["İş Emri", "İşçilik", "Malzeme", "Dış Hizmet", "Toplam Maliyet", "Planlı Bakım Payı %", "Arıza Sayısı", "Plansız Duruş (saat)", "Ort. Onarım Süresi (saat)", "Arızalar Arası Süre (saat)", "İşçilik Saati"]


def _deger(g: dict) -> list:
    return [g["emir"], g["iscilik"], g["malzeme"], g["dis"], g["toplam"], g["planli_pay"], g["ariza"] or None, g["durus"] or None, g["mttr"], g["mtbf"], g["saat"] or None]


def rapor_yaz(cikti: Path, r: dict, planli_hedef: Decimal, calisma: Decimal | None) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    t = r["toplam"]
    _baslik(ws, ["Gösterge", "Değer"], (48, 18))
    for etiket, deger, para_mi in (("İş emri", t["emir"], False), ("İşçilik maliyeti", t["iscilik"], True), ("Malzeme maliyeti", t["malzeme"], True), ("Dış hizmet maliyeti", t["dis"], True), ("Toplam bakım maliyeti", t["toplam"], True),
                                   ("Planlı bakım maliyeti", t["planli"], True), ("Plansız bakım maliyeti", t["plansiz"], True), ("Planlı bakım payı % (maliyet)", t["planli_pay"], True),
                                   ("Planlı bakım payı % (iş emri adedi)", t["planli_emir_pay"], True), ("Planlı bakım payı % (işçilik saati)", t["planli_saat_pay"], True), ("Arıza sayısı", t["ariza"], False),
                                   ("Plansız duruş (saat)", t["durus"], True), ("Planlı bakım duruşu (saat)", t["planli_durus"], True), ("Ortalama onarım süresi, MTTR (saat)", t["mttr"], True),
                                   ("Arızalar arası ortalama süre, MTBF (saat)", t["mtbf"], True)):
        _satir(ws, [etiket, deger], para_s=(2,) if para_mi else ())
    ws.append([])
    ws.append([f"Planlı bakım payı hedefi %{planli_hedef}" + (f" · ekipman başına dönem çalışma saati {calisma}" if calisma is not None else " · çalışma saati verilmediği için MTBF hesaplanmadı")])
    for baslik, alan, ad, anahtar in (("Ekipmanlar", "ekipmanlar", "Ekipman", "ekipman"), ("Bakım Türleri", "turler", "Bakım Türü", "tur"), ("Aylar", "aylar", "Ay", "ay")):
        ws = wb.create_sheet(baslik)
        _baslik(ws, [ad] + (["Maliyet Payı %"] if alan == "ekipmanlar" else []) + BASLIK, (26,) + ((10,) if alan == "ekipmanlar" else ()) + (8, 15, 15, 15, 16, 11, 9, 11, 12, 12, 10), "B2")
        n = 2 if alan == "ekipmanlar" else 1
        for g in r[alan]:
            _satir(ws, [g[anahtar]] + ([g["pay"]] if n == 2 else []) + _deger(g), para_s=[2] * (n - 1) + list(range(n + 2, n + 7)) + [n + 9, n + 10],
                   dolgu={n + 6: None if g["planli_pay"] is None or alan == "turler" else KIRMIZI if g["planli_pay"] < planli_hedef else YESIL})
    ws = wb.create_sheet("İş Emirleri")
    _baslik(ws, ["İş Emri No", "Tarih", "Ekipman", "Bakım Türü", "İşçilik", "Malzeme", "Dış Hizmet", "Toplam", "Duruş (saat)", "İşçilik Saati"], (14, 12, 24, 12, 14, 14, 14, 15, 10, 10))
    for e in r["emirler"]:
        _satir(ws, [e["no"], e["tarih"], e["ekipman"], e["tur"], e["iscilik"], e["malzeme"], e["dis"] or None, e["toplam"], e["durus"] or None, e["saat"] or None], para_s=(5, 6, 7, 8), tarih_s=(2,),
               dolgu={4: SARI if e["tur"] == "Plansız" else None})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, planli_hedef: Decimal = Decimal(70), ariza_esik: int = 3, calisma: Decimal | None = None) -> dict:
    emirler, uy = oku(yol)
    r = raporla(emirler, planli_hedef, ariza_esik, calisma)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, planli_hedef, calisma)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Bakım işçilik, malzeme ve dış hizmet maliyetlerini planlı/plansız bakım oranıyla raporlar.")
    p.add_argument("--emirler", type=Path, help="İş Emri No, Tarih, Ekipman, Bakım Türü, İşçilik Maliyeti, Malzeme Maliyeti, Dış Hizmet Maliyeti, Duruş Süresi (saat), İşçilik Saati")
    p.add_argument("--planli", default="70", help="Planlı bakım maliyet payı hedefi, yüzde; altı uyarılır (varsayılan 70)")
    p.add_argument("--ariza", type=int, default=3, help="Bir ekipmanda bu kadar ve üstü arıza uyarılır (varsayılan 3)")
    p.add_argument("--calisma-saat", help="Dönemde ekipman başına planlı çalışma saati; verilirse arızalar arası süre (MTBF) hesaplanır")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "bakim_maliyet_performans_raporu.xlsx")
    a = p.parse_args(argv)
    yol = a.emirler or ORNEK / "bakim_is_emirleri.csv"
    hedef, calisma = para(a.planli), para(a.calisma_saat) if a.calisma_saat else None
    if hedef is None or not 0 <= hedef <= 100 or a.ariza < 1 or (a.calisma_saat and (calisma is None or calisma <= 0)):
        print("[X] --planli 0 ile 100 arasında; --ariza en az 1; --calisma-saat sıfırdan büyük olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, hedef, a.ariza, calisma)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {t['emir']} iş emri · toplam {t['toplam']:,.2f} (işçilik {t['iscilik']:,.2f}, malzeme {t['malzeme']:,.2f}, dış hizmet {t['dis']:,.2f})" + (f" · planlı pay %{t['planli_pay']}" if t["planli_pay"] is not None else "")
          + f" · {t['ariza']} arıza" + (f", ortalama onarım {t['mttr']} saat" if t["mttr"] is not None else ""))
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
