import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def u(kod, baslik, marka="", tip="", renk="", materyal="", ebat="", model="", ham=None):
    return {"kod": kod, "ham": ham if ham is not None else baslik, "baslik": baslik, "marka": marka, "tip": tip, "renk": renk, "materyal": materyal, "ebat": ebat, "model": model}


def veri():
    return [u("A", "Kahve Makinesi Lumora En Ucuz!!! Siyah Siyah", "Lumora", "Kahve Makinesi", "Siyah"), u("B", "Lumora Çay Makinesi 1.7 L Paslanmaz Çelik Gri", "Lumora", "Çay Makinesi", "Gri", "Paslanmaz Çelik"),
            u("C", "TERMOS", "Nordik", "Termos", "Mavi")]


class KontrolTesti(unittest.TestCase):
    def test_sorunlar_ve_oneri(self):
        r = main.kontrol(veri())
        self.assertEqual([(s["kod"], s["uzunluk"], s["sorunlar"], s["oneri"], s["esit"]) for s in r["urunler"]],
                         [("A", 44, ["Marka başta değil", "Ürün tipi markadan önce", "Yasaklı ifade: en ucuz", "Özel karakter", "Tekrar eden kelime: siyah"], "Lumora Kahve Makinesi Siyah", False),
                          ("B", 45, [], "Lumora Çay Makinesi 1.7 L Paslanmaz Çelik Gri", True),
                          ("C", 6, ["Kısa başlık", "Marka yok", "Renk yok", "Büyük harf"], "Nordik Termos Mavi", False)])
        self.assertEqual(r["ozet"]["Yasaklı ifade"], 1)
        self.assertEqual(len(r["ozet"]), 9)
        self.assertEqual(r["uyarilar"], [])

    def test_secenekler(self):
        r = main.kontrol(veri()[:1], en_az=10, en_cok=40, yasakli=["siyah"])
        self.assertEqual((r["urunler"][0]["sorunlar"], r["urunler"][0]["oneri"]),
                         (["Uzun başlık", "Marka başta değil", "Ürün tipi markadan önce", "Yasaklı ifade: siyah", "Özel karakter", "Tekrar eden kelime: siyah"], "Lumora Kahve Makinesi En Ucuz Siyah"))
        r = main.kontrol([u("X", "Lumora Su Isıtıcı Gri 1.7 Litre Paslanmaz", "Lumora", "Su Isıtıcı", ham="Lumora  Su Isıtıcı Gri 1.7 Litre Paslanmaz "),
                          u("Y", "lumora su ısıtıcı gri 1.7 litre paslanmaz", model="SI-17")])
        self.assertEqual([(s["sorunlar"], s["ayni"], s["oneri"]) for s in r["urunler"]],
                         [(["Fazla boşluk", "Aynı başlık"], ["Y"], "Lumora Su Isıtıcı Gri 1.7 Litre Paslanmaz"), (["Model yok", "Aynı başlık"], ["X"], "lumora su ısıtıcı gri 1.7 litre paslanmaz SI-17")])
        self.assertEqual([x["tur"] for x in main.kontrol([u("Z", "Bir başlık")])["uyarilar"]], ["Marka sütunu yok"])
        self.assertEqual([main.tr_baslik(k) for k in ("İSTANBUL", "ILIK", "çay", "")], ["İstanbul", "Ilık", "Çay", ""])
        with self.assertRaises(ValueError):
            main.kontrol([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "urun_basliklari.csv", cikti)
            self.assertEqual([(s["kod"], s["sorunlar"], None if s["esit"] else s["oneri"]) for s in r["urunler"]],
                             [("EV-101", [], None),
                              ("EV-102", ["Marka başta değil", "Ürün tipi markadan önce", "Renk yok", "Yasaklı ifade: en ucuz", "Özel karakter"], "Lumora Su Isıtıcı Paslanmaz Çelik Gri"),
                              ("EV-103", ["Marka başta değil", "Ürün tipi markadan önce", "Renk yok", "Büyük harf"], "Lumora Tost Makinesi Yapışmaz Plaka Kırmızı"),
                              ("TK-201", ["Marka yok"], "Evora Nevresim Takımı Pamuklu Çift Kişilik Beyaz"),
                              ("TK-202", ["Renk yok", "Materyal yok", "Yasaklı ifade: ucretsiz kargo"], "Evora Banyo Havlusu Havlu Seti 4'lü Lacivert Pamuk"),
                              ("KZ-301", ["Kısa başlık", "Marka yok", "Renk yok", "Materyal yok"], "Nordik Saklama Kabı Şeffaf Cam"),
                              ("KZ-302", ["Kısa başlık", "Fazla boşluk"], "Nordik Döküm Tava 28 cm Siyah"),
                              ("KZ-303", ["Uzun başlık", "Tekrar eden kelime: termos"], "Nordik Çelik Termos 750 ml Mavi Sıcak Soğuk Tutan Bardak Kahve Çay Su İçin Kamp Ofis Seyahat Uyumlu Hediye"),
                              ("KZ-304", ["Renk yok", "Aynı başlık"], "Nordik Çelik Termos 750 ml Mavi Yeşil"), ("KZ-305", ["Aynı başlık"], None)])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Başlık Kontrolü", "Sorun Özeti", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, z = Path(tmp) / "u.csv", Path(tmp) / "y.csv"
            y.write_text("SKU;Ürün Adı;Marka\nA;  Bir  başlık ;M\nB;;M\n;x;M\n", encoding="utf-8")
            z.write_text("Yasaklı İfade\nOrijinal\n\nGarantili Ürün\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["kod"], x["ham"], x["baslik"], x["marka"], x["tip"]) for x in l], [x["tur"] for x in uy]), ([("A", "  Bir  başlık ", "Bir başlık", "M", "")], ["Başlık boş"]))
            self.assertEqual(main.yasakli_oku(z), ["orijinal", "garantili urun"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--urunler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--min", "5", "--max", "60", "--yasakli", str(z), "--ek-yasakli", "bir;"]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--min", "0"], ["--min", "50", "--max", "40"])], [2] * 2)
            self.assertEqual((main.main(["--urunler", str(Path(tmp) / "yok.csv")]), main.main(arg + ["--yasakli", str(Path(tmp) / "yok.csv")])), (1, 1))
            y.write_text("SKU;Marka\nA;M\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
