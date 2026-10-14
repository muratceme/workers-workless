"""
Günlük Giriş-Çıkış Listesi — Workers / Workless kod bloğu
Ön Büro › Resepsiyonist
Rezervasyon listesinden seçilen günün gelecek (arrival), ayrılacak (departure) ve konaklamaya devam eden misafirlerini çıkarır.
  Gelecek       = giriş tarihi seçilen gün olan, iptal / no-show olmayan rezervasyonlar
  Ayrılacak     = çıkış tarihi seçilen gün olan rezervasyonlar
  Devam eden    = giriş < gün < çıkış
  Bakiye        = toplam tutar − ödenen; ödeme şekli acente / cari / fatura ise misafirden tahsil edilmez
  Gece sonu dolu oda = gelecek + devam eden (aynı odadaki rezervasyonlar bir kez sayılır)
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
SUTUN = {"no": ("rezervasyon no", "rez no", "voucher", "rezervasyon"), "misafir": ("misafir", "misafir adi", "ad soyad", "konuk"), "giris": ("giris tarihi", "giris", "check in", "gelis"),
         "cikis": ("cikis tarihi", "cikis", "check out", "ayrilis"), "oda": ("oda no", "oda numarasi", "oda"), "tip": ("oda tipi", "oda turu"), "yetiskin": ("yetiskin", "kisi", "kisi sayisi", "pax"),
         "cocuk": ("cocuk", "cocuk sayisi"), "kanal": ("kanal", "acente", "kaynak"), "pansiyon": ("pansiyon", "konaklama tipi", "board"), "istek": ("ozel istek", "not", "notlar", "istek"),
         "durum": ("durum", "rezervasyon durumu", "statu"), "tutar": ("toplam tutar", "tutar", "konaklama tutari"), "odenen": ("odenen", "on odeme", "tahsil edilen", "depozito"),
         "odeme": ("odeme sekli", "odeme", "odeme tipi")}
CARI = ("acente", "cari", "fatura", "kurum", "voucher", "sirket")
SIFIR = Decimal(0)


def _durum(x) -> str | None:
    k = katla(x)
    if not k or k in ("kesin", "onayli", "garantili", "garanti", "aktif", "rezerve"):
        return "Kesin"
    if "opsiyon" in k or k in ("tentative", "beklemede"):
        return "Opsiyon"
    if "iptal" in k or "no show" in k or k in ("noshow", "cancelled", "canceled", "gelmedi"):
        return "İptal"
    if k in ("konakliyor", "in house", "giris yapti", "check in", "icerde"):
        return "Konaklıyor"
    if k in ("cikis yapti", "check out", "checked out", "ayrildi"):
        return "Çıkış yaptı"
    return None


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("giris", "cikis")):
        no = metin(r["no"])
        if not no and not metin(r["giris"]):
            continue
        g, c, durum = tarih(r["giris"]), tarih(r["cikis"]), _durum(r["durum"])
        say = {a: (para(r[a]) if metin(r[a]) else None) for a in ("yetiskin", "cocuk", "tutar", "odenen")}
        kim = no or metin(r["misafir"]) or f"Satır {r['_satir']}"
        if g is None or c is None or durum is None or any(metin(r[a]) and v is None for a, v in say.items()):
            uy.append(uyari("Yüksek", "Satır okunamadı", kim, f"Satır {r['_satir']}: giriş / çıkış tarihi, durum ('{metin(r['durum'])}') ya da sayılardan biri okunamadı. Rezervasyon listeye alınmadı."))
            continue
        if c <= g:
            uy.append(uyari("Yüksek", "Çıkış girişten önce", kim, f"Giriş {g:%d.%m.%Y}, çıkış {c:%d.%m.%Y}. Rezervasyon listeye alınmadı."))
            continue
        if no and katla(no) in gorulen:
            uy.append(uyari("Orta", "Rezervasyon iki kez yazılmış", kim, f"Satır {r['_satir']}: ikinci satır alınmadı."))
            continue
        if no:
            gorulen.add(katla(no))
        tutar, odenen = say["tutar"], say["odenen"] or SIFIR
        cari = any(k in katla(r["odeme"]) for k in CARI)
        l.append({"no": no, "misafir": metin(r["misafir"]), "giris": g, "cikis": c, "gece": (c - g).days, "oda": metin(r["oda"]), "tip": metin(r["tip"]) or "Belirtilmemiş",
                  "yetiskin": int(say["yetiskin"]) if say["yetiskin"] is not None else None, "cocuk": int(say["cocuk"] or 0), "kanal": metin(r["kanal"]), "pansiyon": metin(r["pansiyon"]),
                  "istek": metin(r["istek"]), "durum": durum, "tutar": tutar, "odenen": odenen, "odeme": metin(r["odeme"]), "cari": cari,
                  "bakiye": None if tutar is None else tutar - odenen})
    return l, uy


def _odeme_durumu(x: dict) -> str:
    if x["cari"]:
        return "Acente / cari"
    if x["bakiye"] is None:
        return "Tutar yok"
    if x["bakiye"] <= 0:
        return "Ödendi"
    return "Kısmi ödeme" if x["odenen"] > 0 else "Ödenmedi"


def analiz(rez: list[dict], gun: date, oda_sayisi: int | None = None) -> dict:
    aktif = [x for x in rez if x["durum"] != "İptal"]
    gelen = sorted((x for x in aktif if x["giris"] == gun), key=lambda x: (x["oda"] == "", katla(x["oda"]), katla(x["misafir"])))
    giden = sorted((x for x in aktif if x["cikis"] == gun), key=lambda x: (katla(x["oda"]), katla(x["misafir"])))
    devam = sorted((x for x in aktif if x["giris"] < gun < x["cikis"]), key=lambda x: (katla(x["oda"]), katla(x["misafir"])))
    for x in gelen + giden + devam:
        x["odeme_durumu"] = _odeme_durumu(x)
    uy = []
    gece = gelen + devam
    odalar = {}
    for x in gece:
        if x["oda"]:
            odalar.setdefault(katla(x["oda"]), []).append(x)
    for l in odalar.values():
        if len(l) > 1:
            uy.append(uyari("Yüksek", "Oda çakışması", f"Oda {l[0]['oda']}", f"{len(l)} rezervasyon aynı gece aynı odada: {', '.join(x['no'] or x['misafir'] for x in l)}."))
    giden_oda = {katla(x["oda"]) for x in giden if x["oda"]}
    for x in gelen:
        kim = f"{x['no']} {x['misafir']}".strip()
        if not x["oda"]:
            uy.append(uyari("Orta", "Oda atanmamış", kim, f"{x['tip']} · {x['gece']} gece. Gelişten önce oda atayın."))
        elif katla(x["oda"]) in giden_oda:
            uy.append(uyari("Bilgi", "Oda aynı gün boşalıyor", kim, f"Oda {x['oda']} bugün çıkış yapacak misafirden boşalıyor; erken girişte oda hazır olmayabilir (kat hizmetlerine öncelik bildirin)."))
        if x["durum"] == "Opsiyon":
            uy.append(uyari("Orta", "Opsiyon kesinleşmemiş", kim, "Bugün gelecek rezervasyon hâlâ opsiyonlu; garanti / ön ödeme durumunu kontrol edin."))
        if x["durum"] in ("Konaklıyor", "Çıkış yaptı"):
            uy.append(uyari("Bilgi", "Giriş zaten yapılmış", kim, f"Durum '{x['durum']}'."))
    for x in giden:
        if x["odeme_durumu"] in ("Ödenmedi", "Kısmi ödeme") and x["durum"] != "Çıkış yaptı":
            uy.append(uyari("Bilgi", "Çıkışta tahsilat", f"{x['no']} {x['misafir']}".strip(), f"Oda {x['oda'] or '-'}: {x['bakiye']:,.2f} bakiye çıkışta tahsil edilecek (ekstralar hariç)."))
    for x in devam:
        if x["durum"] in ("Kesin", "Opsiyon"):
            uy.append(uyari("Orta", "Giriş kaydı yok", f"{x['no']} {x['misafir']}".strip(), f"Giriş {x['giris']:%d.%m.%Y} geçmiş ama durum '{x['durum']}'; no-show mu, giriş mi işlenmedi kontrol edin."))
    dolu = len({katla(x["oda"]) for x in gece if x["oda"]}) + sum(1 for x in gece if not x["oda"])
    kisi = sum((x["yetiskin"] or 0) + x["cocuk"] for x in gece)
    ozet = {"gun": gun, "gelen": len(gelen), "giden": len(giden), "devam": len(devam), "dolu": dolu, "kisi": kisi,
            "doluluk": yuvarla(Decimal(dolu) * 100 / oda_sayisi) if oda_sayisi else None, "istek": sum(1 for x in gelen if x["istek"]),
            "tahsil": sum((x["bakiye"] for x in giden if x["odeme_durumu"] in ("Ödenmedi", "Kısmi ödeme")), SIFIR)}
    if oda_sayisi and dolu > oda_sayisi:
        uy.append(uyari("Yüksek", "Oda sayısı aşıldı", f"{dolu} dolu oda", f"Gece sonu dolu oda {dolu}, otelin oda sayısı {oda_sayisi} (overbooking)."))
    return {"ozet": ozet, "gelen": gelen, "giden": giden, "devam": devam, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    o = r["ozet"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    _baslik(ws, ["Gösterge", "Değer"], (30, 18))
    for a, d in (("Tarih", o["gun"]), ("Gelecek", o["gelen"]), ("Ayrılacak", o["giden"]), ("Devam eden", o["devam"]), ("Gece sonu dolu oda", o["dolu"]),
                 ("Gece sonu kişi", o["kisi"]), ("Doluluk %", o["doluluk"]), ("Özel istekli geliş", o["istek"]), ("Çıkışta tahsil edilecek", o["tahsil"])):
        _satir(ws, [a, d], para_s=(2,) if isinstance(d, Decimal) else (), tarih_s=(2,) if isinstance(d, date) else ())
    ws = wb.create_sheet("Gelecekler")
    _baslik(ws, ["Oda", "Misafir", "Rez. No", "Oda Tipi", "Yetişkin", "Çocuk", "Gece", "Çıkış", "Pansiyon", "Kanal", "Durum", "Özel İstek", "Ödeme", "Bakiye"],
            (8, 24, 12, 14, 8, 7, 6, 11, 10, 14, 10, 34, 14, 12))
    for x in r["gelen"]:
        _satir(ws, [x["oda"] or "Atanmamış", x["misafir"], x["no"], x["tip"], x["yetiskin"], x["cocuk"], x["gece"], x["cikis"], x["pansiyon"], x["kanal"], x["durum"], x["istek"],
                    x["odeme_durumu"], x["bakiye"]], para_s=(14,), tarih_s=(8,), sar=(12,),
               dolgu={1: SARI if not x["oda"] else None, 12: MAVI if x["istek"] else None, 11: SARI if x["durum"] == "Opsiyon" else None})
    ws = wb.create_sheet("Ayrılacaklar")
    _baslik(ws, ["Oda", "Misafir", "Rez. No", "Oda Tipi", "Giriş", "Gece", "Kanal", "Ödeme", "Toplam", "Ödenen", "Bakiye"], (8, 24, 12, 14, 11, 6, 14, 14, 12, 12, 12))
    for x in r["giden"]:
        _satir(ws, [x["oda"], x["misafir"], x["no"], x["tip"], x["giris"], x["gece"], x["kanal"], x["odeme_durumu"], x["tutar"], x["odenen"], x["bakiye"]], para_s=(9, 10, 11), tarih_s=(5,),
               dolgu={8: SARI if x["odeme_durumu"] in ("Ödenmedi", "Kısmi ödeme") else None})
    ws = wb.create_sheet("Devam Edenler")
    _baslik(ws, ["Oda", "Misafir", "Rez. No", "Oda Tipi", "Giriş", "Çıkış", "Durum"], (8, 24, 12, 14, 11, 11, 12))
    for x in r["devam"]:
        _satir(ws, [x["oda"], x["misafir"], x["no"], x["tip"], x["giris"], x["cikis"], x["durum"]], tarih_s=(5, 6))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, gun: date, oda_sayisi: int | None = None) -> dict:
    rez, uy = oku(yol)
    if not rez:
        raise ValueError("Okunabilen rezervasyon yok.")
    r = analiz(rez, gun, oda_sayisi)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Seçilen günün gelecek, ayrılacak ve devam eden misafir listesini çıkarır.")
    p.add_argument("--rezervasyonlar", type=Path, help="Rezervasyon No, Misafir, Giriş, Çıkış, Oda No, Oda Tipi, Yetişkin, Çocuk, Kanal, Pansiyon, Özel İstek, Durum, Toplam Tutar, Ödenen, Ödeme Şekli")
    p.add_argument("--tarih", help="Liste günü (GG.AA.YYYY; varsayılan bugün, örnek veride 15.10.2026)")
    p.add_argument("--oda-sayisi", type=int, help="Otelin satılabilir oda sayısı (doluluk yüzdesi için)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "giris_cikis_listesi.xlsx")
    a = p.parse_args(argv)
    yol = a.rezervasyonlar or ORNEK / "rezervasyonlar.csv"
    gun = tarih(a.tarih) if a.tarih else (date(2026, 10, 15) if a.rezervasyonlar is None else date.today())
    if gun is None or (a.oda_sayisi is not None and a.oda_sayisi < 1):
        print("[X] --tarih GG.AA.YYYY olmalı; --oda-sayisi en az 1")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, gun, a.oda_sayisi)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {gun:%d.%m.%Y}: {o['gelen']} gelecek · {o['giden']} ayrılacak · {o['devam']} devam eden · gece sonu {o['dolu']} dolu oda"
          + (f" (%{o['doluluk']})" if o["doluluk"] is not None else ""))
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Liste: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
