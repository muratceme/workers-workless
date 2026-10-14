import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
GUN = date(2026, 10, 15)


class OrnekTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "rezervasyonlar.csv", cikti, GUN, 20)
            o = r["ozet"]
            self.assertEqual((o["gelen"], o["giden"], o["devam"], o["dolu"], o["kisi"], o["doluluk"], o["istek"], o["tahsil"]), (8, 6, 7, 14, 30, D("70.00"), 7, D(23100)))
            self.assertEqual([x["no"] for x in r["gelen"]], ["R-7105", "R-7101", "R-7106", "R-7102", "R-7107", "R-7104", "R-7103", "R-7108"])
            self.assertEqual({x["no"]: x["odeme_durumu"] for x in r["giden"]},
                             {"R-7051": "Ödenmedi", "R-7052": "Ödendi", "R-7053": "Acente / cari", "R-7054": "Kısmi ödeme", "R-7055": "Kısmi ödeme", "R-7056": "Acente / cari"})
            self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"] if u["onem"] != "Bilgi"),
                             [("Orta", "Giriş kaydı yok", "R-7035 Tolga Erdem"), ("Orta", "Oda atanmamış", "R-7103 Elif Şahin"), ("Orta", "Oda atanmamış", "R-7108 Selin Aydın"),
                              ("Orta", "Opsiyon kesinleşmemiş", "R-7105 Can Yıldız"), ("Yüksek", "Oda çakışması", "Oda 205")])
            self.assertEqual(sorted(u["kim"] for u in r["uyarilar"] if u["tur"] == "Çıkışta tahsilat"), ["R-7051 Ali Çelik", "R-7054 Gizem Ekinci", "R-7055 Hakan Polat"])
            self.assertEqual([u["kim"] for u in r["uyarilar"] if u["tur"] == "Oda aynı gün boşalıyor"], ["R-7105 Can Yıldız"])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Gelecekler", "Ayrılacaklar", "Devam Edenler", "Uyarılar"])

    def test_baska_gun_ve_asim(self):
        rez, _ = main.oku(ORNEK / "rezervasyonlar.csv")
        r = main.analiz(rez, date(2026, 10, 16), 5)
        self.assertEqual((r["ozet"]["gelen"], r["ozet"]["giden"]), (1, 4))
        self.assertIn("Oda sayısı aşıldı", [u["tur"] for u in r["uyarilar"]])


class OkumaTesti(unittest.TestCase):
    def test_durum(self):
        self.assertEqual([main._durum(s) for s in ("", "Garantili", "opsiyonlu", "No Show", "In House", "Checked Out", "x")],
                         ["Kesin", "Kesin", "Opsiyon", "İptal", "Konaklıyor", "Çıkış yaptı", None])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "r.csv"
            y.write_text("Rez No;Konuk;Check In;Check Out;Oda;Pax;Durum;Tutar;Ödenen;Ödeme\n1;A;01.10.2026;03.10.2026;101;2;;1.000;;Kredi kartı\n1;A;01.10.2026;03.10.2026;101;2;;;;\n"
                         "2;B;02.10.2026;01.10.2026;;1;;;;\n3;C;x;03.10.2026;;1;;;;\n4;D;01.10.2026;02.10.2026;;1;belki;;;\n5;E;01.10.2026;04.10.2026;102;1;Kesin;500;500;Acente\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["gece"], x["bakiye"], x["cari"]) for x in l], [u["tur"] for u in uy]),
                             ([("1", 2, D(1000), False), ("5", 3, D(0), True)], ["Rezervasyon iki kez yazılmış", "Çıkış girişten önce", "Satır okunamadı", "Satır okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            self.assertEqual(main.main(["--rezervasyonlar", str(y), "--tarih", "01.10.2026", "--oda-sayisi", "10", "--cikti", c]), 0)
            self.assertEqual([main.main(["--rezervasyonlar", str(y), "--tarih", v, "--cikti", c]) for v in ("x", "32.10.2026")], [2, 2])
            self.assertEqual(main.main(["--oda-sayisi", "0", "--cikti", c]), 2)
            self.assertEqual(main.main(["--rezervasyonlar", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Rez No;Tutar\n1;5\n", encoding="utf-8")
            self.assertEqual(main.main(["--rezervasyonlar", str(y), "--cikti", c]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
