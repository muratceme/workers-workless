"""
Kazanılmış Değer Analizi — Workers / Workless kod bloğu
Planlama ve Proje Kontrol › Planlama Mühendisi

Planlanan değer, kazanılmış değer ve gerçekleşen maliyetten SPI, CPI ve tamamlanma maliyet tahminini hesaplar.
  PV = planlanan değer,  EV = kazanılmış değer,  AC = gerçekleşen maliyet   (hepsi kümülatif; dönemsel verilirse kod toplar),  BAC = toplam bütçe
  Program sapması SV = EV − PV;        maliyet sapması CV = EV − AC
  Program performansı SPI = EV ÷ PV;   maliyet performansı CPI = EV ÷ AC
  Tamamlanma tahmini EAC  = BAC ÷ CPI                              (bugünkü maliyet verimi sürerse)
     seçenek 1            = AC + (BAC − EV)                        (kalan iş bütçesinde biterse)
     seçenek 2            = AC + (BAC − EV) ÷ (CPI × SPI)          (maliyet ve program verimi birlikte sürerse)
  Kalan maliyet ETC = EAC − AC;   tamamlanma sapması VAC = BAC − EAC
  Gereken verim TCPI = (BAC − EV) ÷ (BAC − AC)                     (kalan işi kalan bütçeyle bitirmek için gereken maliyet performansı)
  Süre tahmini (kaba)     = planlanan dönem sayısı ÷ SPI
BAC verilmezse planlanan değerlerin toplamı alınır. EV ve AC'si boş dönemler "ileri dönem" sayılır; analiz verisi olan son döneme göre yapılır.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"donem": ("donem", "ay", "tarih", "hafta"), "pv": ("planlanan deger", "pv", "planlanan", "butce"), "ev": ("kazanilmis deger", "kazanilan deger", "ev", "kazanilmis"),
         "ac": ("gerceklesen maliyet", "ac", "gerceklesen", "fiili maliyet")}
Z = Decimal(0)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("donem", "pv")):
        d = metin(r["donem"])[:10] if not isinstance(r["donem"], (date, datetime)) else f"{r['donem']:%d.%m.%Y}"
        if not d:
            continue
        v = {a: para(r[a]) for a in ("pv", "ev", "ac")}
        if v["pv"] is None or v["pv"] < 0:
            raise ValueError(f"{yol.name} satır {r['_satir']}: '{d}' döneminin planlanan değeri okunamadı ('{metin(r['pv'])}').")
        for a, ad in (("ev", "Kazanılmış değer"), ("ac", "Gerçekleşen maliyet")):
            if v[a] is None and metin(r[a]) or (v[a] is not None and v[a] < 0):
                raise ValueError(f"{yol.name} satır {r['_satir']}: '{d}' döneminin {ad.lower()} alanı okunamadı ('{metin(r[a])}').")
        if (v["ev"] is None) != (v["ac"] is None):
            uy.append(uyari("Yüksek", "Dönem verisi eksik", d, "Kazanılmış değer ve gerçekleşen maliyetten yalnız biri girilmiş; eksik olan 0 sayıldı."))
            v["ev"], v["ac"] = v["ev"] or Z, v["ac"] or Z
        l.append({"donem": d} | v)
    return l, uy


def analiz(donemler: list[dict], bac: Decimal | None = None, kumulatif: bool = False) -> dict:
    if not donemler:
        raise ValueError("Analiz edilecek dönem yok.")
    gercek = [i for i, d in enumerate(donemler) if d["ev"] is not None]
    if not gercek:
        raise ValueError("Hiçbir dönemde kazanılmış değer ve gerçekleşen maliyet yok.")
    if any(donemler[i]["ev"] is None for i in range(gercek[-1])):
        raise ValueError("Verisi olan dönemlerin arasında boş dönem var; kazanılmış değer ve gerçekleşen maliyet ilk dönemden itibaren kesintisiz girilmelidir.")
    kum, satirlar = {"pv": Z, "ev": Z, "ac": Z}, []
    for d in donemler:
        s = {"donem": d["donem"], "ileri": d["ev"] is None}
        for a in ("pv", "ev", "ac"):
            if d[a] is not None:
                if kumulatif and d[a] < kum[a]:
                    raise ValueError(f"'{d['donem']}' döneminde kümülatif {a.upper()} önceki dönemden küçük. Veriler dönemselse --kumulatif vermeyin.")
                kum[a] = d[a] if kumulatif else kum[a] + d[a]
            s[a] = kum[a] if d[a] is not None else None
        if not s["ileri"]:
            s |= {"sv": s["ev"] - s["pv"], "cv": s["ev"] - s["ac"], "spi": yuvarla(s["ev"] / s["pv"]) if s["pv"] else None, "cpi": yuvarla(s["ev"] / s["ac"]) if s["ac"] else None}
        satirlar.append(s)
    toplam_pv = satirlar[-1]["pv"]
    bac = bac if bac is not None else toplam_pv
    if bac <= 0:
        raise ValueError("Toplam bütçe (BAC) pozitif olmalı.")
    son = satirlar[gercek[-1]]
    pv, ev, ac = son["pv"], son["ev"], son["ac"]
    cpi, spi = (ev / ac if ac else None), (ev / pv if pv else None)
    eac = yuvarla(bac * ac / ev) if ev and ac else None
    ozet = {"donem": son["donem"], "bac": bac, "pv": pv, "ev": ev, "ac": ac, "sv": ev - pv, "cv": ev - ac, "spi": son["spi"], "cpi": son["cpi"], "tamamlanma": yuvarla(ev / bac * 100), "planlanan": yuvarla(pv / bac * 100),
            "eac": eac, "etc": eac - ac if eac is not None else None, "vac": bac - eac if eac is not None else None, "eac_butce": ac + (bac - ev),
            "eac_ikili": yuvarla(ac + (bac - ev) / (cpi * spi)) if cpi and spi else None, "tcpi": yuvarla((bac - ev) / (bac - ac)) if bac > ac else None,
            "plan_donem": len(donemler), "gecen_donem": len(gercek), "sure_tahmini": yuvarla(Decimal(len(donemler)) / spi) if spi else None,
            "program": "—" if son["spi"] is None else "Programın gerisinde" if ev < pv else "Programın ilerisinde" if ev > pv else "Programda",
            "maliyet": "—" if son["cpi"] is None else "Bütçenin üzerinde" if ev < ac else "Bütçenin altında" if ev > ac else "Bütçede"}
    return {"donemler": satirlar, "ozet": ozet, "bac_farki": bac != toplam_pv, "toplam_pv": toplam_pv}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    o = r["ozet"]
    ws = wb.active
    ws.title = "Özet"
    _baslik(ws, ["Gösterge", "Değer", "Açıklama"], (44, 18, 90))
    kirmizi = {"Program sapması (SV)": o["sv"] < 0, "Maliyet sapması (CV)": o["cv"] < 0, "Tamamlanma sapması (VAC)": o["vac"] is not None and o["vac"] < 0}
    for ad, d, ac in (("Analiz dönemi", o["donem"], f"{o['gecen_donem']}. dönem / planlanan {o['plan_donem']} dönem"), ("Toplam bütçe (BAC)", o["bac"], ""), ("Planlanan değer (PV)", o["pv"], f"Planlanan ilerleme %{o['planlanan']}"),
                      ("Kazanılmış değer (EV)", o["ev"], f"Gerçekleşen ilerleme %{o['tamamlanma']}"), ("Gerçekleşen maliyet (AC)", o["ac"], ""), ("Program sapması (SV)", o["sv"], o["program"]),
                      ("Maliyet sapması (CV)", o["cv"], o["maliyet"]), ("Program performansı (SPI)", o["spi"], "1'in altı: planlanandan yavaş"), ("Maliyet performansı (CPI)", o["cpi"], "1'in altı: harcanan her 1 birim 1'den az iş üretiyor"),
                      ("Tamamlanma tahmini (EAC)", o["eac"], "BAC ÷ CPI: bugünkü maliyet verimi sürerse"), ("Kalan maliyet (ETC)", o["etc"], "EAC − AC"), ("Tamamlanma sapması (VAC)", o["vac"], "BAC − EAC; eksi: bütçe aşımı"),
                      ("EAC, kalan iş bütçesinde biterse", o["eac_butce"], "AC + (BAC − EV)"), ("EAC, maliyet ve program verimi sürerse", o["eac_ikili"], "AC + (BAC − EV) ÷ (CPI × SPI)"),
                      ("Gereken verim (TCPI)", o["tcpi"], "Kalan işi kalan bütçeyle bitirmek için gereken CPI; 1'in çok üzerindeyse bütçe gerçekçi değildir"),
                      ("Süre tahmini (dönem, kaba)", o["sure_tahmini"], f"Planlanan {o['plan_donem']} dönem ÷ SPI; iş programı analizinin yerini tutmaz")):
        _satir(ws, [ad, d if d is not None else "hesaplanamadı", ac], para_s=(2,), dolgu={2: KIRMIZI if kirmizi.get(ad) else None}, sar=(3,))
    ws = wb.create_sheet("Dönem Bazında")
    _baslik(ws, ["Dönem", "PV (kümülatif)", "EV (kümülatif)", "AC (kümülatif)", "SV", "CV", "SPI", "CPI"], (14, 16, 16, 16, 15, 15, 8, 8))
    for s in r["donemler"]:
        _satir(ws, [s["donem"], s["pv"], s["ev"], s["ac"], s.get("sv"), s.get("cv"), s.get("spi"), s.get("cpi")], para_s=(2, 3, 4, 5, 6, 7, 8),
               dolgu={5: KIRMIZI if s.get("sv") is not None and s["sv"] < 0 else None, 6: KIRMIZI if s.get("cv") is not None and s["cv"] < 0 else None, 1: MAVI if s["ileri"] else None})
    ws.append([])
    ws.append(["Mavi işaretli dönemler ileri dönemlerdir (yalnız plan). Tüm değerler kümülatiftir."])
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bac: Decimal | None = None, kumulatif: bool = False) -> dict:
    donemler, uy = oku(yol)
    r = analiz(donemler, bac, kumulatif)
    o = r["ozet"]
    if r["bac_farki"]:
        uy.append(uyari("Orta", "BAC planlanan değer toplamından farklı", yol.name, f"Verilen BAC {o['bac']:,.2f}, planlanan değerlerin toplamı {r['toplam_pv']:,.2f}. İleri dönemlerin planı dosyada eksik olabilir."))
    if o["ev"] > o["bac"]:
        uy.append(uyari("Yüksek", "Kazanılmış değer bütçeyi aşıyor", o["donem"], f"EV {o['ev']:,.2f} > BAC {o['bac']:,.2f}. Kazanılmış değer bütçe fiyatlarıyla ölçülmelidir."))
    if o["tcpi"] is None:
        uy.append(uyari("Yüksek", "Bütçe tükendi", o["donem"], "Gerçekleşen maliyet toplam bütçeye ulaştı veya aştı; TCPI hesaplanamaz, bütçe revizyonu gerekir."))
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="PV, EV ve AC'den SPI, CPI ve tamamlanma maliyet tahminini hesaplar.")
    p.add_argument("--donemler", type=Path, default=ORNEK / "donemler.csv", help="Dönem, Planlanan Değer, Kazanılmış Değer, Gerçekleşen Maliyet (ileri dönemlerde son ikisi boş)")
    p.add_argument("--bac", help="Toplam bütçe (verilmezse planlanan değerlerin toplamı)")
    p.add_argument("--kumulatif", action="store_true", help="Dosyadaki değerler kümülatifse verin (varsayılan: dönemsel)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kazanilmis_deger.xlsx")
    a = p.parse_args(argv)
    bac = para(a.bac) if a.bac else None
    if a.bac and (bac is None or bac <= 0):
        print("[X] --bac pozitif sayı olmalı")
        return 2
    if not a.donemler.exists():
        print(f"[X] Dosya bulunamadı: {a.donemler}")
        return 1
    try:
        r = calistir(a.donemler, a.cikti, bac, a.kumulatif)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {o['donem']} · BAC {o['bac']:,.2f} · PV {o['pv']:,.2f} · EV {o['ev']:,.2f} · AC {o['ac']:,.2f}")
    print(f"[{'!' if o['sv'] < 0 or o['cv'] < 0 else 'OK'}] SPI {o['spi'] if o['spi'] is not None else '—'} ({o['program']}) · CPI {o['cpi'] if o['cpi'] is not None else '—'} ({o['maliyet']})")
    if o["eac"] is not None:
        print(f"[{'!' if o['vac'] < 0 else 'OK'}] Tamamlanma tahmini {o['eac']:,.2f} · bütçeye göre {o['vac']:+,.2f}" + (f" · gereken verim TCPI {o['tcpi']}" if o["tcpi"] is not None else ""))
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
