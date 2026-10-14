import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 14)


def f(no, durum="Açık", ilk=(10, 1), test=None, teklif=None, tutar=None, son=None, plan=None, satis=None, neden="", danisman="A", model="M"):
    t = {k: (date(2026, *v) if v else None) for k, v in (("ilk", ilk), ("test", test), ("teklif", teklif), ("son", son), ("plan", plan), ("satis", satis))}
    sn = max(v for k, v in t.items() if k in ("ilk", "test", "teklif", "son") and v)
    return {"no": no, "musteri": f"Müşteri {no}", "telefon": "", "danisman": danisman, "model": model, **t, "son_temas": sn, "tutar": None if tutar is None else D(tutar),
            "durum": durum, "neden": neden, "takas": False, "kredi": False}


class AnalizTesti(unittest.TestCase):
    def test_oncelik(self):
        l = [f("1", test=(10, 2), teklif=(10, 3), tutar=1_500_000, son=(10, 8)), f("2", test=(10, 9)), f("3", plan=(10, 14), son=(10, 13)), f("4", teklif=(10, 12), tutar=2_000_000),
             f("5", ilk=(8, 20), plan=(9, 1)), f("6", ilk=(10, 12), test=(10, 13), plan=(10, 20)),
             f("7", "Satış", ilk=(9, 1), test=(9, 2), satis=(9, 21), danisman="B"), f("8", "Kayıp", ilk=(9, 1), test=(9, 3), neden="Fiyat", danisman="B"), f("9", "Kayıp", ilk=(9, 5), danisman="B"),
             f("10", "Satış", ilk=(9, 1))]
        r = main.analiz(l, BUGUN)
        self.assertEqual([(x["no"], x["oncelik"], x["neden_oncelik"]) for x in r["acik"]],
                         [("1", 1, "Teklif verildi, 6 gündür temas yok"), ("5", 1, "Planlanan arama 43 gün gecikti"), ("3", 1, "Planlanan arama bugün"),
                          ("2", 2, "Test sürüşü yapıldı, teklif yok, 5 gündür temas yok"), ("4", 3, "Takip"), ("6", 3, "Planlı arama 20.10.2026")])
        t = r["toplam"]
        self.assertEqual((t["firsat"], t["test"], t["teklif"], t["satis"], t["kayip"], t["acik"], t["kapanis"], t["test_kapanis"], t["sure"], t["oncelik1"]),
                         (10, 5, 2, 2, 2, 6, D("50.00"), D("50.00"), D("20.0"), 3))
        self.assertEqual([(s["ad"], s["kapanis"]) for s in r["danisman"]], [("A", D("100.00")), ("B", D("33.33"))])
        self.assertEqual(r["nedenler"], [{"neden": "Belirtilmemiş", "adet": 1, "pay": D("50.00")}, {"neden": "Fiyat", "adet": 1, "pay": D("50.00")}])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Bilgi", "Fırsat soğuyor", "5 Müşteri 5"), ("Bilgi", "Kayıp nedeni yok", "9 Müşteri 9"), ("Bilgi", "Satış tarihi yok", "10 Müşteri 10"),
                          ("Orta", "Planlanan arama kaçırıldı", "5 Müşteri 5")])
        self.assertEqual([x["oncelik"] for x in main.analiz(l, BUGUN, takip=7)["acik"]][:4], [1, 1, 3, 3])
        with self.assertRaises(ValueError):
            main.analiz([], BUGUN)

    def test_durum(self):
        self.assertEqual([main._durum(x) for x in ("", "Sıcak", "Satıldı", "Rakip", "x")], ["Açık", "Açık", "Satış", "Kayıp", None])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "firsatlar.csv", cikti, BUGUN)
            t = r["toplam"]
            self.assertEqual((t["firsat"], t["satis"], t["kayip"], t["acik"], t["kapanis"], t["test_kapanis"], t["oncelik1"]), (160, 54, 44, 62, D("55.10"), D("68.85"), 30))
            self.assertEqual(r["acik"][0]["oncelik"], 1)
            self.assertEqual(load_workbook(cikti).sheetnames, ["Geri Dönüş Listesi", "Danışman", "Model", "Kayıp Nedenleri", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "f.csv"
            y.write_text("Müşteri;Temsilci;Gelis Tarihi;Test Sürüşü;Teklif Tarihi;Teklif Tutarı;Sonuç\nA;X;01.10.2026;03.10.2026;;;\nB;X;01.10.2026;Evet;;;\nC;X;01.10.2026;;;1.000.000;belki\n"
                         "D;;01.10.2026;;05.10.2026;1.200.000;Satıldı\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["musteri"], x["danisman"], x["son_temas"], x["tutar"], x["durum"]) for x in l], len(uy)),
                             ([("A", "X", date(2026, 10, 3), None, "Açık"), ("D", "Atanmamış", date(2026, 10, 5), D(1200000), "Satış")], 2))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--firsatlar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "14.10.2026", "--takip", "2", "--soguma", "20"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--takip", "0"], ["--soguma", "3"])], [2] * 3)
            self.assertEqual(main.main(["--firsatlar", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Müşteri;Sonuç\nA;belki\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
