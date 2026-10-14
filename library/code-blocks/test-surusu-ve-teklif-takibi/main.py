"""
Test Sürüşü ve Teklif Takibi — Workers / Workless kod bloğu
Satış › Satış Danışmanı
Test sürüşü yapan ve teklif verilen müşterileri izler, açık fırsatlar için geri dönüş önceliği belirler, dönüşüm oranlarını çıkarır.
  Öncelik 1 (bugün ara) = planlanan arama tarihi bugün ya da geçmiş, veya teklif verilmiş ve son temastan --takip (3) gün geçmiş
  Öncelik 2             = test sürüşü yapılmış, teklif verilmemiş ve son temastan --takip gün geçmiş
  Öncelik 3             = diğer açık fırsatlar;  soğuyor = son temastan --soguma (30) günden fazla geçmiş
  Kapanış oranı %       = satış ÷ (satış + kayıp) × 100;  test sürüşü / teklif oranı = ilgili adım ÷ fırsat × 100
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
SUTUN = {"no": ("kayit no", "firsat no", "no"), "musteri": ("musteri", "ad soyad", "musteri adi"), "telefon": ("telefon", "gsm", "cep telefonu"), "danisman": ("danisman", "satis danismani", "temsilci"),
         "model": ("model", "ilgilenilen model", "arac"), "ilk": ("ilk temas tarihi", "ilk temas", "kayit tarihi", "gelis tarihi"), "test": ("test surusu tarihi", "test surusu"),
         "teklif": ("teklif tarihi",), "tutar": ("teklif tutari", "tutar", "teklif"), "son": ("son temas tarihi", "son temas", "son gorusme tarihi", "son gorusme"),
         "plan": ("planlanan arama", "sonraki arama", "sonraki temas", "takip tarihi"), "durum": ("durum", "sonuc"), "neden": ("kayip nedeni", "neden"),
         "takas": ("takas",), "kredi": ("kredi", "finansman"), "satis": ("satis tarihi", "fatura tarihi")}
TARIHLER = ("ilk", "test", "teklif", "son", "plan", "satis")
SIFIR = Decimal(0)


def _durum(x) -> str | None:
    k = katla(x)
    if not k or k in ("acik", "takipte", "devam", "sicak", "ilik", "soguk", "beklemede"):
        return "Açık"
    if k in ("satis", "satildi", "kazanildi", "siparis", "fatura", "teslim"):
        return "Satış"
    if k in ("kayip", "kaybedildi", "vazgecti", "olumsuz", "iptal", "rakip"):
        return "Kayıp"
    return None


def _evet(x) -> bool:
    return katla(x) in ("evet", "e", "var", "x", "1")


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("musteri", "durum")):
        if not metin(r["musteri"]):
            continue
        kim = metin(r["no"]) or metin(r["musteri"])
        t = {a: (tarih(r[a]) if metin(r[a]) else None) for a in TARIHLER}
        tutar = para(r["tutar"]) if metin(r["tutar"]) else None
        d = _durum(r["durum"])
        if d is None or any(metin(r[a]) and t[a] is None for a in TARIHLER) or (metin(r["tutar"]) and tutar is None):
            uy.append(uyari("Yüksek", "Satır okunamadı", kim, f"Satır {r['_satir']}: durum ('{metin(r['durum'])}'), tarihlerden biri ya da teklif tutarı okunamadı."))
            continue
        son = max((v for a, v in t.items() if a in ("ilk", "test", "teklif", "son") and v), default=None)
        l.append({"no": metin(r["no"]), "musteri": metin(r["musteri"]), "telefon": metin(r["telefon"]), "danisman": metin(r["danisman"]) or "Atanmamış", "model": metin(r["model"]) or "Belirtilmemiş",
                  **t, "son_temas": son, "tutar": tutar, "durum": d, "neden": metin(r["neden"]), "takas": _evet(r["takas"]), "kredi": _evet(r["kredi"])})
    return l, uy


def _oran(a: int, b: int) -> Decimal | None:
    return yuvarla(Decimal(a) * 100 / b) if b else None


def analiz(firsatlar: list[dict], bugun: date, takip: int = 3, soguma: int = 30) -> dict:
    if not firsatlar:
        raise ValueError("Okunabilen fırsat yok.")
    uy, acik = [], []
    for f in firsatlar:
        f["gun"] = (bugun - f["son_temas"]).days if f["son_temas"] else None
        f["oncelik"], f["neden_oncelik"] = None, ""
        kim = f"{f['no']} {f['musteri']}".strip()
        if f["durum"] == "Satış" and not f["satis"]:
            uy.append(uyari("Bilgi", "Satış tarihi yok", kim, "Satış olarak işaretli ama satış tarihi yazılmamış; satış süresi hesabına girmedi."))
        if f["durum"] == "Kayıp" and not f["neden"]:
            uy.append(uyari("Bilgi", "Kayıp nedeni yok", kim, "Kayıp fırsatta neden yazılmamış."))
        if f["durum"] != "Açık":
            continue
        bekliyor = f["gun"] is not None and f["gun"] >= takip
        if f["plan"] and f["plan"] <= bugun:
            f["oncelik"], f["neden_oncelik"] = 1, "Planlanan arama " + ("bugün" if f["plan"] == bugun else f"{(bugun - f['plan']).days} gün gecikti")
        elif f["teklif"] and bekliyor:
            f["oncelik"], f["neden_oncelik"] = 1, f"Teklif verildi, {f['gun']} gündür temas yok"
        elif f["test"] and not f["teklif"] and bekliyor:
            f["oncelik"], f["neden_oncelik"] = 2, f"Test sürüşü yapıldı, teklif yok, {f['gun']} gündür temas yok"
        else:
            f["oncelik"], f["neden_oncelik"] = 3, ("Planlı arama " + f"{f['plan']:%d.%m.%Y}" if f["plan"] else "Takip")
        if f["plan"] and (bugun - f["plan"]).days > takip:
            uy.append(uyari("Orta", "Planlanan arama kaçırıldı", kim, f"{f['danisman']}: arama {f['plan']:%d.%m.%Y} için planlanmıştı ({(bugun - f['plan']).days} gün önce)."))
        if f["gun"] is not None and f["gun"] > soguma:
            uy.append(uyari("Bilgi", "Fırsat soğuyor", kim, f"{f['danisman']} · {f['model']}: {f['gun']} gündür temas yok; kapatın ya da yeniden arayın."))
        acik.append(f)
    acik.sort(key=lambda f: (f["oncelik"], -(f["tutar"] or SIFIR), -(f["gun"] or 0), katla(f["musteri"])))

    def ozet(fn):
        g = {}
        for f in firsatlar:
            g.setdefault(fn(f), []).append(f)
        satir = []
        for ad, l in g.items():
            satis = [f for f in l if f["durum"] == "Satış"]
            kayip = [f for f in l if f["durum"] == "Kayıp"]
            test = [f for f in l if f["test"]]
            sure = [(f["satis"] - f["ilk"]).days for f in satis if f["satis"] and f["ilk"] and f["satis"] >= f["ilk"]]
            satir.append({"ad": ad, "firsat": len(l), "test": len(test), "teklif": sum(1 for f in l if f["teklif"]), "satis": len(satis), "kayip": len(kayip),
                          "acik": sum(1 for f in l if f["durum"] == "Açık"), "test_oran": _oran(len(test), len(l)), "teklif_oran": _oran(sum(1 for f in l if f["teklif"]), len(l)),
                          "kapanis": _oran(len(satis), len(satis) + len(kayip)),
                          "test_kapanis": _oran(sum(1 for f in test if f["durum"] == "Satış"), sum(1 for f in test if f["durum"] != "Açık")),
                          "sure": yuvarla(Decimal(sum(sure)) / len(sure), "0.1") if sure else None,
                          "oncelik1": sum(1 for f in l if f["oncelik"] == 1)})
        return sorted(satir, key=lambda s: (-s["firsat"], katla(s["ad"])))
    nedenler = {}
    for f in firsatlar:
        if f["durum"] == "Kayıp":
            nedenler[f["neden"] or "Belirtilmemiş"] = nedenler.get(f["neden"] or "Belirtilmemiş", 0) + 1
    kayip = sum(nedenler.values())
    return {"acik": acik, "danisman": ozet(lambda f: f["danisman"]), "model": ozet(lambda f: f["model"]), "toplam": ozet(lambda f: "TOPLAM")[0],
            "nedenler": [{"neden": n, "adet": a, "pay": _oran(a, kayip)} for n, a in sorted(nedenler.items(), key=lambda x: (-x[1], katla(x[0])))], "uyarilar": uy}


OZET_BASLIK = ["Fırsat", "Test Sürüşü", "Teklif", "Satış", "Kayıp", "Açık", "Test Sürüşü %", "Teklif %", "Kapanış %", "Test Sürüşü Yapanlarda Kapanış %", "Ort. Satış Süresi (gün)", "Bugün Aranacak"]


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Geri Dönüş Listesi"
    _baslik(ws, ["Öncelik", "Neden", "Kayıt No", "Müşteri", "Telefon", "Danışman", "Model", "Test Sürüşü", "Teklif Tarihi", "Teklif Tutarı", "Son Temas", "Gün", "Planlanan Arama", "Takas", "Kredi"],
            (7, 36, 9, 20, 14, 16, 16, 11, 11, 13, 11, 6, 11, 6, 6), "C2")
    for f in r["acik"]:
        _satir(ws, [f["oncelik"], f["neden_oncelik"], f["no"], f["musteri"], f["telefon"], f["danisman"], f["model"], f["test"], f["teklif"], f["tutar"], f["son_temas"], f["gun"], f["plan"],
                    "Evet" if f["takas"] else "", "Evet" if f["kredi"] else ""], para_s=(10,), tarih_s=(8, 9, 11, 13), sar=(2,),
               dolgu={1: KIRMIZI if f["oncelik"] == 1 else SARI if f["oncelik"] == 2 else None})
    for ad, l in (("Danışman", r["danisman"]), ("Model", r["model"])):
        ws = wb.create_sheet(ad)
        _baslik(ws, [ad] + OZET_BASLIK, (20,) + (9,) * 12)
        for s in l + [r["toplam"]]:
            n = _satir(ws, [s["ad"], s["firsat"], s["test"], s["teklif"], s["satis"], s["kayip"], s["acik"], s["test_oran"], s["teklif_oran"], s["kapanis"], s["test_kapanis"], s["sure"],
                            s["oncelik1"]], para_s=(8, 9, 10, 11))
            if s["ad"] == "TOPLAM":
                for h in ws[n]:
                    h.font = Font(bold=True)
    ws = wb.create_sheet("Kayıp Nedenleri")
    _baslik(ws, ["Neden", "Adet", "Pay %"], (36, 8, 8))
    for n in r["nedenler"]:
        _satir(ws, [n["neden"], n["adet"], n["pay"]], para_s=(3,))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, takip: int = 3, soguma: int = 30) -> dict:
    firsatlar, uy = oku(yol)
    r = analiz(firsatlar, bugun, takip, soguma)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Test sürüşü ve teklif verilen müşterileri izleyip geri dönüş önceliği belirler.")
    p.add_argument("--firsatlar", type=Path, help="Kayıt No, Müşteri, Telefon, Danışman, Model, İlk Temas, Test Sürüşü Tarihi, Teklif Tarihi, Teklif Tutarı, Son Temas, Planlanan Arama, "
                                                  "Durum (Açık / Satış / Kayıp), Kayıp Nedeni, Takas, Kredi, Satış Tarihi")
    p.add_argument("--tarih", help="Rapor günü GG.AA.YYYY (varsayılan bugün; örnek veride 14.10.2026)")
    p.add_argument("--takip", type=int, default=3, help="Son temastan kaç gün sonra geri dönülmeli (varsayılan 3)")
    p.add_argument("--soguma", type=int, default=30, help="Kaç gün temas olmazsa fırsat soğuyor sayılır (varsayılan 30)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "test_surusu_teklif_takibi.xlsx")
    a = p.parse_args(argv)
    yol = a.firsatlar or ORNEK / "firsatlar.csv"
    bugun = tarih(a.tarih) if a.tarih else (date(2026, 10, 14) if a.firsatlar is None else date.today())
    if bugun is None or a.takip < 1 or a.soguma <= a.takip:
        print("[X] --tarih GG.AA.YYYY; --takip en az 1, --soguma --takip'ten büyük olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, bugun, a.takip, a.soguma)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {t['firsat']} fırsat · {t['satis']} satış · {t['kayip']} kayıp · {t['acik']} açık · kapanış %{t['kapanis']} · test sürüşü yapanlarda %{t['test_kapanis']}")
    print(f"     Bugün aranacak (öncelik 1): {t['oncelik1']} · öncelik 2: {sum(1 for f in r['acik'] if f['oncelik'] == 2)}")
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Liste: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
