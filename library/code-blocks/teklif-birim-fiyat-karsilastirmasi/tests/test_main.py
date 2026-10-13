import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def t(poz, miktar, firma, fiyat):
    return {"poz": poz, "tanim": "", "birim": "m2", "miktar": D(miktar), "firma": firma, "fiyat": D(fiyat), "satir": 0}


TEKLIF = lambda: [t("P1", 100, "A", 50), t("P1", 100, "B", 60), t("P1", 100, "C", 55), t("P2", 10, "A", 200), t("P2", 10, "B", 100), t("P2", 10, "C", 210),   # noqa: E731
                  t("P3", 20, "A", 30), t("P3", 20, "B", 28), t("P4", 5, "A", 400)]


class AnalizTesti(unittest.TestCase):
    def test_karsilastirma(self):
        r = main.analiz(TEKLIF())
        self.assertEqual([(p["poz"], p["en_firma"], p["en_fiyat"], p["ortalama"], p["medyan"], p["fark"], p["teklif"], p["tutar"], p["isaret"]) for p in r["pozlar"]],
                         [("P1", "A", D(50), D("55.00"), D("55.00"), D("9.09"), 3, D("5000.00"), {}), ("P2", "B", D(100), D("170.00"), D("200.00"), D("41.18"), 3, D("1000.00"), {"b": "düşük"}),
                          ("P3", "B", D(28), D("29.00"), D("29.00"), D("3.45"), 2, D("560.00"), {}), ("P4", "A", D(400), D("400.00"), D("400.00"), D("0.00"), 1, D("2000.00"), {})])
        self.assertEqual([(f["firma"], f["tutar"], f["poz"], f["eksik"], f["tam"], f.get("sira"), f["karma_poz"], f["karma_tutar"]) for f in r["firmalar"]],
                         [("A", D("9600.00"), 4, [], True, 1, 2, D("7000.00")), ("B", D("7560.00"), 3, ["P4"], False, None, 2, D("1560.00")), ("C", D("7600.00"), 2, ["P3", "P4"], False, None, 0, D(0))])
        o = r["ozet"]
        self.assertEqual((o["poz"], o["firma"], o["karma"], o["en_tam_firma"], o["en_tam"], o["tasarruf"]), (4, 3, D("8560.00"), "A", D("9600.00"), D("1040.00")))
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]), [("Aşırı düşük fiyat", "P2 · B"), ("Eksik teklif", "B"), ("Eksik teklif", "C"), ("Tek teklif", "P4")])

    def test_esikler_esitlik_ve_tutarsizlik(self):
        f = lambda *fy, **kw: main.analiz([t("P", 1, chr(65 + i), v) for i, v in enumerate(fy)], **kw)["pozlar"][0]   # noqa: E731
        self.assertEqual((f(70, 100, 100)["isaret"], f("69.99", 100, 100)["isaret"], f(100, 100, 130)["isaret"], f(100, 100, "130.01")["isaret"]), ({}, {"a": "düşük"}, {}, {"c": "yüksek"}))
        self.assertEqual((f(50, 100)["isaret"], f(80, 100, 100, dusuk=D(10))["isaret"], f(10, 20, 30, 40)["medyan"], f(40, 40, 50)["en_firma"]), ({}, {"a": "düşük"}, D("25.00"), "A"))
        r = main.analiz([t("P1", 10, "Beta", 5), t("p1", 12, "Alfa", 5), t("P1", 10, "BETA", 4), t("P1", 10, "Gama", 6)])
        self.assertEqual(([x["firma"] for x in r["firmalar"]], r["pozlar"][0]["en_firma"], r["pozlar"][0]["miktar"], r["ozet"]["karma"], sorted(u["tur"] for u in r["uyarilar"])),
                         (["Alfa", "Beta", "Gama"], "Alfa", D(10), D("50.00"), ["Miktar tutarsız", "Mükerrer teklif"]))
        r = main.analiz([t("P1", 1, "A", 5), t("P2", 1, "B", 5)])
        self.assertEqual((r["ozet"]["en_tam"], r["ozet"]["tasarruf"], [x.get("sira") for x in r["firmalar"]]), (None, None, [None, None]))
        r = main.analiz([t("P1", 1, "A", 5), t("P1", 1, "B", 4), t("P2", 2, "A", 3), t("P2", 2, "B", 6)])
        self.assertEqual(([(x["firma"], x["tutar"], x["sira"]) for x in r["firmalar"]], r["ozet"]["karma"], r["ozet"]["tasarruf"]), ([("A", D("11.00"), 1), ("B", D("16.00"), 2)], D("10.00"), D("1.00")))
        with self.assertRaises(ValueError):
            main.analiz([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "teklifler.csv", cikti)
            o = r["ozet"]
            self.assertEqual([(p["poz"], p["en_firma"], p["tutar"]) for p in r["pozlar"]], [("P-01", "Örnek Yapı Ltd.", D("444000.00")), ("P-02", "Deneme İnşaat A.Ş.", D("678500.00")),
                                                                                           ("P-03", "Deneme İnşaat A.Ş.", D("301600.00")), ("P-04", "Örnek Yapı Ltd.", D("336000.00")), ("P-05", "Deneme İnşaat A.Ş.", D("46800.00"))])
            self.assertEqual([(f["firma"], f["tutar"], f["tam"], f.get("sira")) for f in r["firmalar"]],
                             [("Deneme İnşaat A.Ş.", D("1886900.00"), True, 1), ("Numune Taahhüt Ltd.", D("1655100.00"), False, None), ("Örnek Yapı Ltd.", D("2058750.00"), True, 2)])
            self.assertEqual((o["karma"], o["en_tam_firma"], o["tasarruf"], [(u["tur"], u["kim"]) for u in r["uyarilar"]]),
                             (D("1806900.00"), "Deneme İnşaat A.Ş.", D("80000.00"), [("Aşırı düşük fiyat", "P-03 · Deneme İnşaat A.Ş."), ("Eksik teklif", "Numune Taahhüt Ltd.")]))
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Karşılaştırma", "Firma Toplamları", "Karma Cetvel", "Uyarılar"])
            self.assertEqual([c.value for c in wb["Karşılaştırma"][2]][:8], ["P-01", "Alçı sıva yapılması", "m2", 2400, 210, 198, 185, 185])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "t.csv"
            y.write_text("Poz;Tanım;Birim;Metraj;Taşeron;Fiyat\nP1;Sıva;m2;12,5;A;1.250,50\nP1;Sıva;m2;12.5;B;1300\nP1;Sıva;m2;12,5;C;\nP2;Boya;m2;0;A;10\nP2;Boya;m2;5;;10\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["poz"], x["firma"], x["miktar"], x["fiyat"]) for x in l], [u["tur"] for u in uy]), ([("P1", "A", D("12.5"), D("1250.50")), ("P1", "B", D("12.5"), D(1300))], ["Teklif satırı okunamadı"] * 2))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--teklifler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(arg + ["--dusuk", "100"]), 2)
            self.assertEqual(main.main(arg + ["--yuksek", "x"]), 2)
            self.assertEqual(main.main(["--teklifler", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Poz;Firma\nP1;A\n", "Poz;Miktar;Firma;Fiyat\nP1;x;A;y\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
