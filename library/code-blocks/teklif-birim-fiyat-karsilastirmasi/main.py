"""
Teklif Birim Fiyat Karşılaştırması — Gerisi Bende kod bloğu
İhale ve Teklif › İhale Mühendisi

Taşeron ve tedarikçi tekliflerini poz bazında karşılaştırır; firma toplamlarını ve en düşük fiyatlardan oluşan karma teklif cetvelini çıkarır.
  Poz tutarı         = miktar × birim fiyat
  En düşük           = pozdaki en düşük birim fiyat (eşitlikte adı alfabetik önce gelen firma)
  Medyan             = pozdaki tekliflerin ortancası;  ortalamadan fark % = (ortalama − en düşük) ÷ ortalama × 100
  Aşırı düşük/yüksek = en az 3 teklif varsa, medyanın --dusuk (%30) altındaki veya --yuksek (%30) üstündeki fiyat
  Firma toplamı      = Σ poz tutarı;  yalnız bütün pozlara fiyat veren firmalar ("tam teklif") toplamda sıralanır
  Karma cetvel       = her pozda en düşük fiyatı veren firma;  karma toplam = Σ miktar × en düşük fiyat
  Karma tasarruf     = en düşük tam teklif − karma toplam
Fiyatlar aynı para biriminde, aynı kapsamda (KDV, nakliye, montaj dahil veya hariç) olmalıdır; kod kapsam farkını bilemez. En düşük fiyat en uygun teklif demek değildir.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"poz": ("poz no", "poz", "kalem no", "sira no"), "tanim": ("is kalemi", "tanim", "aciklama", "imalat", "is tanimi"), "birim": ("birim", "olcu birimi"), "miktar": ("miktar", "metraj"),
         "firma": ("firma", "teklif veren", "taseron", "tedarikci"), "fiyat": ("birim fiyat", "teklif birim fiyati", "fiyat", "teklif")}
Z = Decimal(0)


def olcu(x) -> Decimal | None:
    """Miktar: 12,5 = 12.5; 1.250,5 = bin iki yüz elli buçuk."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    else:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("poz", "miktar", "firma", "fiyat")):
        poz, firma = metin(r["poz"]), metin(r["firma"])
        if not poz and not firma:
            continue
        kim = f"{poz} · {firma} (satır {r['_satir']})"
        if not metin(r["fiyat"]) and poz and firma:
            continue                                                                             # fiyat verilmemiş poz: eksik teklif olarak sayılır
        m, f = olcu(r["miktar"]), para(r["fiyat"])
        if not poz or not firma or m is None or m <= 0 or f is None or f <= 0:
            uy.append(uyari("Yüksek", "Teklif satırı okunamadı", kim, f"Miktar '{metin(r['miktar'])}', birim fiyat '{metin(r['fiyat'])}'. Poz ve firma dolu, miktar ve fiyat pozitif olmalı. Satır alınmadı."))
            continue
        l.append({"poz": poz, "tanim": metin(r["tanim"]), "birim": metin(r["birim"]), "miktar": m, "firma": firma, "fiyat": f, "satir": r["_satir"]})
    return l, uy


def medyan(d: list[Decimal]) -> Decimal:
    s, n = sorted(d), len(d)
    return s[n // 2] if n % 2 else (s[n // 2 - 1] + s[n // 2]) / 2


def analiz(teklifler: list[dict], dusuk: Decimal = Decimal(30), yuksek: Decimal = Decimal(30)) -> dict:
    if not teklifler:
        raise ValueError("Karşılaştırılacak teklif yok.")
    uy, pozlar, adlar = [], {}, {}
    for t in teklifler:
        fa = katla(t["firma"])
        adlar.setdefault(fa, t["firma"])
        p = pozlar.setdefault(katla(t["poz"]), {"poz": t["poz"], "tanim": t["tanim"], "birim": t["birim"], "miktar": t["miktar"], "fiyat": {}})
        p["tanim"], p["birim"] = p["tanim"] or t["tanim"], p["birim"] or t["birim"]
        if t["miktar"] != p["miktar"]:
            uy.append(uyari("Orta", "Miktar tutarsız", f"{t['poz']} · {t['firma']}", f"Bu satırda miktar {t['miktar']:f}, pozun ilk satırında {p['miktar']:f}. Karşılaştırmada {p['miktar']:f} kullanıldı; firmalar aynı metraja mı fiyat verdi, kontrol edin."))
        if fa in p["fiyat"]:
            uy.append(uyari("Orta", "Mükerrer teklif", f"{t['poz']} · {t['firma']}", f"Aynı firma bu poza birden fazla fiyat vermiş ({p['fiyat'][fa]:f} ve {t['fiyat']:f}). İlk fiyat kullanıldı."))
            continue
        p["fiyat"][fa] = t["fiyat"]
    firmalar = sorted(adlar, key=lambda f: adlar[f])
    toplam = {f: {"firma": adlar[f], "tutar": Z, "poz": 0, "eksik": [], "karma_poz": 0, "karma_tutar": Z} for f in firmalar}
    liste, karma = [], Z
    for p in pozlar.values():
        fy = p["fiyat"]
        en = min(fy, key=lambda f: (fy[f], adlar[f]))
        ort, med, n = sum(fy.values()) / len(fy), medyan(list(fy.values())), len(fy)
        for f in firmalar:
            if f in fy:
                toplam[f]["tutar"] += p["miktar"] * fy[f]
                toplam[f]["poz"] += 1
            else:
                toplam[f]["eksik"].append(p["poz"])
        isaret = {}
        if n >= 3:
            for f, v in fy.items():
                if v < med * (1 - dusuk / 100):
                    isaret[f] = "düşük"
                    uy.append(uyari("Orta", "Aşırı düşük fiyat", f"{p['poz']} · {adlar[f]}", f"Birim fiyat {v:,.2f}, pozun medyanı {med:,.2f} (%{yuvarla((med - v) / med * 100)} altında). Kapsam, birim veya yazım hatası olabilir; firmadan teyit alın."))
                elif v > med * (1 + yuksek / 100):
                    isaret[f] = "yüksek"
                    uy.append(uyari("Bilgi", "Aşırı yüksek fiyat", f"{p['poz']} · {adlar[f]}", f"Birim fiyat {v:,.2f}, pozun medyanı {med:,.2f} (%{yuvarla((v - med) / med * 100)} üstünde)."))
        elif n == 1:
            uy.append(uyari("Bilgi", "Tek teklif", p["poz"], f"Bu poza yalnız {adlar[en]} fiyat vermiş; karşılaştırma yapılamadı."))
        tutar = yuvarla(p["miktar"] * fy[en])
        karma += tutar
        toplam[en]["karma_poz"] += 1
        toplam[en]["karma_tutar"] += tutar
        liste.append({"poz": p["poz"], "tanim": p["tanim"], "birim": p["birim"], "miktar": p["miktar"], "fiyat": fy, "isaret": isaret, "teklif": n, "en": en, "en_firma": adlar[en], "en_fiyat": fy[en],
                      "ortalama": yuvarla(ort), "medyan": yuvarla(med), "fark": yuvarla((ort - fy[en]) / ort * 100), "tutar": tutar})
    for f, t in toplam.items():
        t["tutar"], t["tam"] = yuvarla(t["tutar"]), not t["eksik"]
        if t["eksik"]:
            uy.append(uyari("Orta", "Eksik teklif", t["firma"], f"{len(t['eksik'])} poza fiyat verilmemiş ({', '.join(t['eksik'][:8])}{'…' if len(t['eksik']) > 8 else ''}). Firma toplamı diğerleriyle karşılaştırılamaz."))
    tam = sorted((t for t in toplam.values() if t["tam"]), key=lambda t: (t["tutar"], t["firma"]))
    for i, t in enumerate(tam, 1):
        t["sira"] = i
    en_tam = tam[0] if tam else None
    ozet = {"poz": len(liste), "firma": len(firmalar), "karma": karma, "en_tam_firma": en_tam["firma"] if en_tam else None, "en_tam": en_tam["tutar"] if en_tam else None,
            "tasarruf": en_tam["tutar"] - karma if en_tam else None, "ayar": (dusuk, yuksek)}
    return {"pozlar": liste, "firmalar": [toplam[f] for f in firmalar], "firma_sira": firmalar, "ozet": ozet, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb, o = Workbook(), r["ozet"]
    ws = wb.active
    ws.title = "Karşılaştırma"
    adlar = [t["firma"] for t in r["firmalar"]]
    nf = len(adlar)
    _baslik(ws, ["Poz No", "İş Kalemi", "Birim", "Miktar"] + adlar + ["En Düşük", "En Düşük Firma", "Ortalama", "Medyan", "Ortalamadan Fark %", "Teklif Sayısı"], (12, 36, 8, 11) + (15,) * nf + (13, 22, 13, 13, 11, 8), "E2")
    for p in r["pozlar"]:
        _satir(ws, [p["poz"], p["tanim"], p["birim"], p["miktar"]] + [p["fiyat"].get(f) for f in r["firma_sira"]] + [p["en_fiyat"], p["en_firma"], p["ortalama"], p["medyan"], p["fark"], p["teklif"]],
               para_s=tuple(range(4, 5 + nf + 1)) + (nf + 7, nf + 8, nf + 9),
               dolgu={5 + j: YESIL if f == p["en"] else SARI if p["isaret"].get(f) == "düşük" else KIRMIZI if p["isaret"].get(f) == "yüksek" else None for j, f in enumerate(r["firma_sira"])})
    ws.append([])
    ws.append(["Yeşil: pozun en düşük fiyatı. Kırmızı: medyanın çok üstünde. Sarı: medyanın çok altında (en düşük fiyat aşırı düşükse Uyarılar sayfasına bakın). Boş hücre: fiyat verilmemiş."])
    ws = wb.create_sheet("Firma Toplamları")
    _baslik(ws, ["Firma", "Teklif Verilen Poz", "Eksik Poz", "Teklif Toplamı", "Tam Teklif", "Sıra", "Karma Cetvelde Poz", "Karma Cetvelde Tutar"], (28, 10, 10, 18, 10, 7, 12, 18))
    for t in r["firmalar"]:
        _satir(ws, [t["firma"], t["poz"], len(t["eksik"]) or None, t["tutar"], "Evet" if t["tam"] else "Hayır", t.get("sira"), t["karma_poz"] or None, t["karma_tutar"] or None], para_s=(4, 8),
               dolgu={5: None if t["tam"] else SARI, 6: YESIL if t.get("sira") == 1 else None})
    ws.append([])
    for ad, deger in (("Karma cetvel toplamı (her pozda en düşük fiyat)", o["karma"]), (f"En düşük tam teklif ({o['en_tam_firma'] or 'yok'})", o["en_tam"]), ("Karma cetvelin tam teklife göre tasarrufu", o["tasarruf"])):
        n = _satir(ws, [ad, None, None, deger], para_s=(4,))
        ws.cell(n, 1).font = Font(bold=True)
    ws = wb.create_sheet("Karma Cetvel")
    _baslik(ws, ["Poz No", "İş Kalemi", "Birim", "Miktar", "Firma", "Birim Fiyat", "Tutar"], (12, 36, 8, 11, 26, 14, 16))
    for p in r["pozlar"]:
        _satir(ws, [p["poz"], p["tanim"], p["birim"], p["miktar"], p["en_firma"], p["en_fiyat"], p["tutar"]], para_s=(4, 6, 7))
    n = _satir(ws, ["Toplam", "", "", None, "", None, o["karma"]], para_s=(7,))
    for c in ws[n]:
        c.font = Font(bold=True)
    ws.append([])
    ws.append(["Karma cetvel bir pazarlık aracıdır. İşi pozlara bölmek her zaman mümkün değildir (mobilizasyon, sorumluluk, asgari iş hacmi); firmalar tek poz için aynı fiyatı vermeyebilir."])
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, **ayar) -> dict:
    teklifler, uy = oku(yol)
    r = analiz(teklifler, **ayar)
    uy += r["uyarilar"]
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Teklifleri poz bazında karşılaştırır, firma toplamlarını ve karma teklif cetvelini çıkarır.")
    p.add_argument("--teklifler", type=Path, default=ORNEK / "teklifler.csv", help="Poz No, İş Kalemi, Birim, Miktar, Firma, Birim Fiyat (her firma-poz bir satır)")
    p.add_argument("--dusuk", default="30", help="Medyanın bu %%'den fazla altındaki fiyat 'aşırı düşük' (varsayılan 30)")
    p.add_argument("--yuksek", default="30", help="Medyanın bu %%'den fazla üstündeki fiyat 'aşırı yüksek' (varsayılan 30)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "teklif_karsilastirma.xlsx")
    a = p.parse_args(argv)
    d, y = olcu(a.dusuk), olcu(a.yuksek)
    if d is None or y is None or not 0 < d < 100 or y <= 0:
        print("[X] --dusuk 0–100 arasında, --yuksek pozitif sayı olmalı")
        return 2
    if not a.teklifler.exists():
        print(f"[X] Dosya bulunamadı: {a.teklifler}")
        return 1
    try:
        r = calistir(a.teklifler, a.cikti, dusuk=d, yuksek=y)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {o['poz']} poz · {o['firma']} firma · karma cetvel toplamı {o['karma']:,.2f}")
    if o["en_tam"] is not None:
        print(f"[OK] En düşük tam teklif: {o['en_tam_firma']} {o['en_tam']:,.2f} · karma cetvel {o['tasarruf']:,.2f} daha düşük")
    else:
        print("[!] Bütün pozlara fiyat veren firma yok; firma toplamları karşılaştırılamaz")
    for u in r["uyarilar"]:
        if u["tur"] in ("Aşırı düşük fiyat", "Eksik teklif"):
            print(f"[!] {u['tur']}: {u['kim']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
