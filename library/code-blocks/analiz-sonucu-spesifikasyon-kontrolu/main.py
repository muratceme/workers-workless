"""
Analiz Sonucu Spesifikasyon Kontrolü — Gerisi Bende kod bloğu
Kalite Kontrol ve Laboratuvar › Kalite Kontrol Teknisyeni

Laboratuvar analiz sonuçlarını ürün spesifikasyon limitleriyle karşılaştırır; uygunsuz partileri ve sınıra yakın değerleri işaretler.
  Uygunsuz          = sonuç < alt limit veya sonuç > üst limit;  nitel parametrede sonuç ≠ beklenen (ör. "Negatif")
  Sınıra yakın      = iki limit varsa: limite uzaklık ≤ (üst − alt) × --yakin %;  tek limit varsa: sonuç, limitin --yakin %'si kadar yakınında   (varsayılan %10)
  "<10" gibi sonuç  = gerçek değer 10'un altında: üst limit ≥ 10 ise uygun, değilse değerlendirilemez;  ">150" → üst limit ≤ 150 ise uygunsuz
  Eksik analiz      = ürünün spesifikasyonunda olan ama partide sonucu bulunmayan parametre
  Parti kararı      = Uygunsuz > Eksik analiz > Değerlendirilemedi > Sınıra yakın > Uygun   (en ağır durum)
Limitler sizin spesifikasyon dosyanızdan gelir; kod mevzuat veya kodeks limiti bilmez. Limite eşit sonuç uygundur.
Ölçüm belirsizliği dikkate alınmaz. Parti serbest bırakma kararı yetkili kişinindir.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SONUC = {"parti": ("parti no", "parti", "lot no", "numune no", "numune"), "urun": ("urun", "urun adi"), "parametre": ("parametre", "analiz", "test", "analiz adi"),
         "sonuc": ("sonuc", "deger", "analiz sonucu"), "tarih": ("analiz tarihi", "tarih")}
SPEC = {"urun": SONUC["urun"], "parametre": SONUC["parametre"], "alt": ("alt limit", "min", "alt sinir", "en az"), "ust": ("ust limit", "max", "ust sinir", "en cok"),
        "beklenen": ("beklenen", "beklenen sonuc", "kabul"), "birim": ("birim",)}
AGIRLIK = {"Uygunsuz": 0, "Eksik analiz": 1, "Değerlendirilemedi": 2, "Sınıra yakın": 3, "Uygun": 4, "Spesifikasyon yok": 5}
DURUM_RENK = {"Uygunsuz": KIRMIZI, "Eksik analiz": KIRMIZI, "Değerlendirilemedi": SARI, "Sınıra yakın": SARI, "Uygun": YESIL}


def olcu(x) -> Decimal | None:
    """Ölçüm: 4,55 = 4.55; 1.250,5 = bin iki yüz elli buçuk. Tek başına nokta ondalık ayırıcıdır."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    else:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yaz(v: Decimal) -> str:
    return f"{v.normalize():f}".replace(".", ",")


def limit_metni(sp: dict) -> str:
    if sp["beklenen"]:
        return f"= {sp['beklenen']}"
    if sp["alt"] is not None and sp["ust"] is not None:
        return f"{yaz(sp['alt'])} – {yaz(sp['ust'])}"
    return f"≥ {yaz(sp['alt'])}" if sp["alt"] is not None else f"≤ {yaz(sp['ust'])}"


def oku_spec(yol: Path) -> tuple[dict, list[dict]]:
    uy, sp = [], {}
    for r in kayitlar(yol, SPEC, ("parametre",)):
        par = metin(r["parametre"])
        if not par:
            continue
        alt, ust, bek = olcu(r["alt"]), olcu(r["ust"]), metin(r["beklenen"])
        kim = f"{metin(r['urun']) or 'Tüm ürünler'} · {par}"
        if (alt is None and metin(r["alt"])) or (ust is None and metin(r["ust"])) or (alt is None and ust is None and not bek) or (alt is not None and ust is not None and alt > ust):
            uy.append(uyari("Yüksek", "Spesifikasyon okunamadı", kim, f"Alt limit '{metin(r['alt'])}', üst limit '{metin(r['ust'])}', beklenen '{bek}'. En az biri dolu olmalı, alt limit üst limitten büyük olamaz. Satır alınmadı."))
            continue
        a = (katla(r["urun"]), katla(par))
        if a in sp:
            uy.append(uyari("Orta", "Mükerrer spesifikasyon", kim, "Aynı ürün ve parametre için birden fazla satır var. İlk satır kullanıldı."))
            continue
        sp[a] = {"urun": metin(r["urun"]), "parametre": par, "alt": alt, "ust": ust, "beklenen": bek, "birim": metin(r["birim"])}
    return sp, uy


def oku_sonuc(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SONUC, ("parti", "parametre", "sonuc")):
        p, par = metin(r["parti"]), metin(r["parametre"])
        if not p and not par:
            continue
        if not p or not par or not metin(r["sonuc"]):
            uy.append(uyari("Yüksek", "Sonuç satırı okunamadı", f"{p or '?'} · {par or '?'} (satır {r['_satir']})", "Parti, parametre veya sonuç boş. Satır alınmadı."))
            continue
        ham = r["sonuc"]
        l.append({"parti": p, "urun": metin(r["urun"]), "parametre": par, "sonuc": yaz(Decimal(str(ham))) if isinstance(ham, (int, float)) else metin(ham), "tarih": tarih(r["tarih"])})
    return l, uy


def degerlendir(ham: str, sp: dict, yakin: Decimal = Decimal(10)) -> tuple[str, str, Decimal | None]:
    """(durum, açıklama, sayısal değer) döndürür."""
    if sp["beklenen"]:
        return ("Uygun", "", None) if katla(ham) == katla(sp["beklenen"]) else ("Uygunsuz", f"Beklenen '{sp['beklenen']}', sonuç '{ham}'.", None)
    m = re.match(r"\s*([<>≤≥])?\s*=?\s*(.*)", ham)
    isaret, v, alt, ust = m.group(1), olcu(m.group(2)), sp["alt"], sp["ust"]
    if v is None:
        return "Değerlendirilemedi", f"Sonuç '{ham}' sayıya çevrilemedi.", None
    yetmez = ("Değerlendirilemedi", f"'{ham}' sonucu limitle ({limit_metni(sp)}) karşılaştırmaya yetmiyor.", None)
    if isaret in ("<", "≤"):
        if alt is not None and v <= alt:
            return "Uygunsuz", f"Sonuç alt limitin ({yaz(alt)}) altında.", None
        return yetmez if alt is not None or (ust is not None and v > ust) else ("Uygun", "", None)
    if isaret in (">", "≥"):
        if ust is not None and v >= ust:
            return "Uygunsuz", f"Sonuç üst limitin ({yaz(ust)}) üzerinde.", None
        return yetmez if ust is not None or (alt is not None and v < alt) else ("Uygun", "", None)
    if alt is not None and v < alt:
        return "Uygunsuz", f"Alt limitin ({yaz(alt)}) {yaz(alt - v)} altında.", v
    if ust is not None and v > ust:
        return "Uygunsuz", f"Üst limitin ({yaz(ust)}) {yaz(v - ust)} üzerinde.", v
    y = yakin / 100
    if alt is not None and ust is not None:
        pay = (ust - alt) * y
        if v - alt <= pay:
            return "Sınıra yakın", f"Alt limite ({yaz(alt)}) {yaz(v - alt)} uzaklıkta.", v
        if ust - v <= pay:
            return "Sınıra yakın", f"Üst limite ({yaz(ust)}) {yaz(ust - v)} uzaklıkta.", v
    elif ust is not None and ust > 0 and v >= ust * (1 - y):
        return "Sınıra yakın", f"Üst limite ({yaz(ust)}) {yaz(ust - v)} uzaklıkta.", v
    elif alt is not None and alt > 0 and v <= alt * (1 + y):
        return "Sınıra yakın", f"Alt limite ({yaz(alt)}) {yaz(v - alt)} uzaklıkta.", v
    return "Uygun", "", v


def analiz(sonuclar: list[dict], spec: dict, yakin: Decimal = Decimal(10)) -> dict:
    if not sonuclar:
        raise ValueError("Değerlendirilecek analiz sonucu yok.")
    if not spec:
        raise ValueError("Spesifikasyon dosyasında geçerli satır yok.")
    uy, satirlar, partiler, par = [], [], {}, {}
    for s in sonuclar:
        uk, pk = katla(s["urun"]), katla(s["parametre"])
        sp = spec.get((uk, pk)) or spec.get(("", pk))
        p = partiler.setdefault(katla(s["parti"]), {"parti": s["parti"], "urun": s["urun"], "uk": uk, "gorulen": set(), "durumlar": [], "tarih": None})
        if uk != p["uk"]:
            uy.append(uyari("Orta", "Ürün adı tutarsız", s["parti"], f"Partide hem '{p['urun']}' hem '{s['urun']}' yazılmış. Her satır kendi ürün adının spesifikasyonuyla değerlendirildi."))
        if pk in p["gorulen"]:
            uy.append(uyari("Bilgi", "Tekrarlı analiz", f"{s['parti']} · {s['parametre']}", "Aynı parametre partide birden fazla kez var; her sonuç ayrı değerlendirildi."))
        p["gorulen"].add(pk)
        if s["tarih"]:
            p["tarih"] = max(p["tarih"] or s["tarih"], s["tarih"])
        if sp is None:
            durum, ac, v, limit, birim = "Spesifikasyon yok", "Bu ürün ve parametre için limit tanımlı değil.", olcu(s["sonuc"]), "", ""
        else:
            durum, ac, v = degerlendir(s["sonuc"], sp, yakin)
            limit, birim = limit_metni(sp), sp["birim"]
            o = par.setdefault(pk, {"parametre": sp["parametre"], "birim": birim, "n": 0, "Uygunsuz": 0, "Sınıra yakın": 0, "Değerlendirilemedi": 0, "degerler": []})
            o["n"] += 1
            if durum in o:
                o[durum] += 1
            if v is not None:
                o["degerler"].append(v)
        p["durumlar"].append(durum)
        satirlar.append(s | {"deger": v, "limit": limit, "birim": birim, "durum": durum, "aciklama": ac})
        if durum in ("Uygunsuz", "Değerlendirilemedi"):
            uy.append(uyari("Yüksek" if durum == "Uygunsuz" else "Orta", durum, f"{s['parti']} · {s['parametre']}", f"Sonuç {s['sonuc']} {birim}".strip() + f", limit {limit}. {ac}"))
        elif durum == "Sınıra yakın":
            uy.append(uyari("Bilgi", durum, f"{s['parti']} · {s['parametre']}", f"Sonuç {s['sonuc']} {birim}".strip() + f", limit {limit}. {ac}"))
    liste = []
    for p in partiler.values():
        beklenen = {pk: sp["parametre"] for (u, pk), sp in spec.items() if u in (p["uk"], "")}
        eksik = sorted(ad for pk, ad in beklenen.items() if pk not in p["gorulen"])
        if eksik:
            uy.append(uyari("Yüksek", "Eksik analiz", p["parti"], f"Spesifikasyonda olup sonucu bulunmayan parametreler: {', '.join(eksik)}."))
        d = p["durumlar"] + (["Eksik analiz"] if eksik else [])
        liste.append({"parti": p["parti"], "urun": p["urun"], "tarih": p["tarih"], "analiz": len(p["durumlar"]), "uygunsuz": d.count("Uygunsuz"), "yakin": d.count("Sınıra yakın"),
                      "belirsiz": d.count("Değerlendirilemedi"), "eksik": eksik, "karar": min(d, key=AGIRLIK.get)})
    ozet = []
    for o in par.values():
        dg = o.pop("degerler")
        ozet.append(o | {"min": min(dg) if dg else None, "max": max(dg) if dg else None, "ort": yuvarla(sum(dg) / len(dg), "0.001") if dg else None})
    say = defaultdict(int)
    for x in liste:
        say[x["karar"]] += 1
    return {"partiler": liste, "sonuclar": satirlar, "parametreler": sorted(ozet, key=lambda o: o["parametre"]), "say": dict(say), "yakin": yakin, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Parti Özeti"
    _baslik(ws, ["Parti No", "Ürün", "Son Analiz", "Analiz Sayısı", "Uygunsuz", "Sınıra Yakın", "Değerlendirilemedi", "Eksik Parametreler", "Karar Önerisi", "Onay / Not (siz doldurun)"], (16, 30, 12, 9, 10, 10, 12, 34, 20, 30))
    for p in r["partiler"]:
        _satir(ws, [p["parti"], p["urun"], p["tarih"], p["analiz"], p["uygunsuz"] or None, p["yakin"] or None, p["belirsiz"] or None, ", ".join(p["eksik"]), p["karar"], ""], tarih_s=(3,),
               dolgu={9: DURUM_RENK.get(p["karar"])}, sar=(8,))
    ws.append([])
    ws.append([f"Sınıra yakın: limite uzaklık, limit aralığının (tek limitte limit değerinin) %{yaz(r['yakin'])}'i içinde. Limite eşit sonuç uygundur. Ölçüm belirsizliği dikkate alınmaz; serbest bırakma kararı yetkili kişinindir."])
    ws = wb.create_sheet("Sonuçlar")
    _baslik(ws, ["Parti No", "Ürün", "Parametre", "Sonuç", "Birim", "Limit", "Durum", "Açıklama", "Analiz Tarihi"], (16, 30, 26, 12, 10, 16, 20, 56, 12))
    for s in r["sonuclar"]:
        _satir(ws, [s["parti"], s["urun"], s["parametre"], s["sonuc"], s["birim"], s["limit"], s["durum"], s["aciklama"], s["tarih"]], tarih_s=(9,), dolgu={7: DURUM_RENK.get(s["durum"])}, sar=(8,))
    ws = wb.create_sheet("Parametre Özeti")
    _baslik(ws, ["Parametre", "Birim", "Sonuç Sayısı", "Uygunsuz", "Sınıra Yakın", "Değerlendirilemedi", "En Düşük", "En Yüksek", "Ortalama"], (26, 10, 9, 10, 10, 12, 12, 12, 12))
    for o in r["parametreler"]:
        _satir(ws, [o["parametre"], o["birim"], o["n"], o["Uygunsuz"] or None, o["Sınıra yakın"] or None, o["Değerlendirilemedi"] or None, o["min"], o["max"], o["ort"]], dolgu={4: KIRMIZI if o["Uygunsuz"] else None})
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(sonuc_yol: Path, spec_yol: Path, cikti: Path, yakin: Decimal = Decimal(10)) -> dict:
    spec, uy = oku_spec(spec_yol)
    sonuclar, uy2 = oku_sonuc(sonuc_yol)
    r = analiz(sonuclar, spec, yakin)
    uy += uy2 + r["uyarilar"]
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Analiz sonuçlarını spesifikasyon limitleriyle karşılaştırır.")
    p.add_argument("--sonuclar", type=Path, default=ORNEK / "analiz_sonuclari.csv", help="Parti No, Ürün, Parametre, Sonuç, Analiz Tarihi")
    p.add_argument("--spesifikasyon", type=Path, help="Ürün (boşsa tüm ürünler), Parametre, Alt Limit, Üst Limit, Beklenen (nitel), Birim")
    p.add_argument("--yakin", default="10", help="Sınıra yakınlık payı, %% (varsayılan 10)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "spesifikasyon_kontrol.xlsx")
    a = p.parse_args(argv)
    spec = a.spesifikasyon or (ORNEK / "spesifikasyon.csv" if a.sonuclar == ORNEK / "analiz_sonuclari.csv" else None)
    y = olcu(a.yakin)
    if y is None or not 0 <= y < 50:
        print("[X] --yakin 0 ile 50 arasında bir sayı olmalı")
        return 2
    if spec is None:
        print("[X] --spesifikasyon dosyası verilmeli")
        return 2
    for yol in (a.sonuclar, spec):
        if not yol.exists():
            print(f"[X] Dosya bulunamadı: {yol}")
            return 1
    try:
        r = calistir(a.sonuclar, spec, a.cikti, y)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    s = r["say"]
    print(f"[OK] {len(r['partiler'])} parti · {len(r['sonuclar'])} sonuç · " + " · ".join(f"{s[k]} {k.lower()}" for k in AGIRLIK if k in s))
    for x in r["partiler"]:
        if x["karar"] in ("Uygunsuz", "Eksik analiz"):
            print(f"[!] {x['parti']} ({x['urun']}): {x['karar']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
