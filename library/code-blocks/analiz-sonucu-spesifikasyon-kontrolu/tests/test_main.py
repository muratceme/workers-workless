import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def sp(alt=None, ust=None, beklenen=""):
    return {"urun": "", "parametre": "X", "alt": None if alt is None else D(alt), "ust": None if ust is None else D(ust), "beklenen": beklenen, "birim": ""}


def d(ham, **kw):
    yakin = kw.pop("yakin", D(10))
    return main.degerlendir(ham, sp(**kw), yakin)[0]


def s(parti, par, sonuc, urun="U"):
    return {"parti": parti, "urun": urun, "parametre": par, "sonuc": sonuc, "tarih": None}


class DegerlendirmeTesti(unittest.TestCase):
    def test_iki_limit(self):
        self.assertEqual([d(v, alt="4.3", ust="4.8") for v in ("4,29", "4,3", "4,35", "4,36", "4,55", "4,74", "4,75", "4,8", "4,81")],
                         ["Uygunsuz", "Sınıra yakın", "Sınıra yakın", "Uygun", "Uygun", "Uygun", "Sınıra yakın", "Sınıra yakın", "Uygunsuz"])
        self.assertEqual((main.degerlendir("4,81", sp(alt="4.3", ust="4.8")), main.degerlendir("4.2", sp(alt="4.3", ust="4.8"))[1], main.degerlendir("4,78", sp(alt="4.3", ust="4.8"))[1]),
                         (("Uygunsuz", "Üst limitin (4,8) 0,01 üzerinde.", D("4.81")), "Alt limitin (4,3) 0,1 altında.", "Üst limite (4,8) 0,02 uzaklıkta."))
        self.assertEqual((d("4,75", alt="4.3", ust="4.8", yakin=D(0)), d("4,8", alt="4.3", ust="4.8", yakin=D(0)), d("5", alt=5, ust=5)), ("Uygun", "Sınıra yakın", "Sınıra yakın"))

    def test_tek_limit_ve_nitel(self):
        self.assertEqual([d(v, ust=100) for v in ("89,9", "90", "100", "100,1", "0")], ["Uygun", "Sınıra yakın", "Sınıra yakın", "Uygunsuz", "Uygun"])
        self.assertEqual([d(v, alt=40) for v in ("39,9", "40", "44", "44,1")], ["Uygunsuz", "Sınıra yakın", "Sınıra yakın", "Uygun"])
        self.assertEqual([d(v, beklenen="Negatif") for v in ("negatif", " NEGATİF ", "Pozitif", "<10")], ["Uygun", "Uygun", "Uygunsuz", "Uygunsuz"])
        self.assertEqual((d("yok", ust=5), d("", ust=5), d("1.250,5", ust=2000), main.limit_metni(sp(alt=40)), main.limit_metni(sp(ust="6.5")), main.limit_metni(sp(alt="4.3", ust="4.80")), main.limit_metni(sp(beklenen="Yok"))),
                         ("Değerlendirilemedi", "Değerlendirilemedi", "Uygun", "≥ 40", "≤ 6,5", "4,3 – 4,8", "= Yok"))

    def test_sansurlu_sonuc(self):
        self.assertEqual([d(v, ust=100) for v in ("<10", "< 100", "<101", ">150", ">100", ">50", "≥100")], ["Uygun", "Uygun", "Değerlendirilemedi", "Uygunsuz", "Uygunsuz", "Değerlendirilemedi", "Uygunsuz"])
        self.assertEqual([d(v, alt=40) for v in ("<40", "<30", "<50", ">40", ">45", ">30")], ["Uygunsuz", "Uygunsuz", "Değerlendirilemedi", "Uygun", "Uygun", "Değerlendirilemedi"])
        self.assertEqual([d(v, alt=10, ust=20) for v in ("<10", "<15", ">20", ">15")], ["Uygunsuz", "Değerlendirilemedi", "Uygunsuz", "Değerlendirilemedi"])


class AnalizTesti(unittest.TestCase):
    def test_parti_karari(self):
        spec = {("u", "ph"): sp(alt="4.3", ust="4.8") | {"parametre": "pH"}, ("u", "tuz"): sp(ust="6.5") | {"parametre": "Tuz", "birim": "%"}, ("", "salmonella"): sp(beklenen="Negatif") | {"parametre": "Salmonella"}}
        tam = lambda p, ph, tuz, sal: [s(p, "pH", ph), s(p, "TUZ", tuz), s(p, "Salmonella", sal)]   # noqa: E731
        r = main.analiz(tam("P1", "4,5", "5", "Negatif") + tam("P2", "4,78", "5", "Negatif") + tam("P3", "4,5", "7", "Negatif") + [s("P4", "pH", "4,5"), s("P4", "Tuz", "6,4")] + tam("P5", "4,5", "<7", "Negatif")
                        + [s("P6", "Salmonella", "Negatif", "Başka"), s("P6", "Nem", "12", "Başka")] + tam("P7", "4,9", "5", "Pozitif") + [s("P7", "pH", "4,6")], spec)
        self.assertEqual([(p["parti"], p["karar"], p["analiz"], p["uygunsuz"], p["yakin"], p["belirsiz"], p["eksik"]) for p in r["partiler"]],
                         [("P1", "Uygun", 3, 0, 0, 0, []), ("P2", "Sınıra yakın", 3, 0, 1, 0, []), ("P3", "Uygunsuz", 3, 1, 0, 0, []), ("P4", "Eksik analiz", 2, 0, 1, 0, ["Salmonella"]),
                          ("P5", "Değerlendirilemedi", 3, 0, 0, 1, []), ("P6", "Uygun", 2, 0, 0, 0, []), ("P7", "Uygunsuz", 4, 2, 0, 0, [])])
        self.assertEqual(r["say"], {"Uygun": 2, "Sınıra yakın": 1, "Uygunsuz": 2, "Eksik analiz": 1, "Değerlendirilemedi": 1})
        self.assertEqual([(o["parametre"], o["n"], o["Uygunsuz"], o["Sınıra yakın"], o["Değerlendirilemedi"], o["min"], o["max"], o["ort"]) for o in r["parametreler"]],
                         [("Salmonella", 6, 1, 0, 0, None, None, None), ("Tuz", 6, 1, 1, 1, D(5), D(7), D("5.680")), ("pH", 7, 1, 1, 0, D("4.5"), D("4.9"), D("4.611"))])
        self.assertEqual([x["durum"] for x in r["sonuclar"] if x["parti"] == "P6"], ["Uygun", "Spesifikasyon yok"])
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Değerlendirilemedi", "P5 · TUZ"), ("Eksik analiz", "P4"), ("Sınıra yakın", "P2 · pH"), ("Sınıra yakın", "P4 · Tuz"), ("Tekrarlı analiz", "P7 · pH"), ("Uygunsuz", "P3 · TUZ"), ("Uygunsuz", "P7 · Salmonella"),
                          ("Uygunsuz", "P7 · pH")])
        for a, b in (([], spec), ([s("P", "pH", "4")], {})):
            with self.assertRaises(ValueError):
                main.analiz(a, b)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "a.xlsx"
            r = main.calistir(ORNEK / "analiz_sonuclari.csv", ORNEK / "spesifikasyon.csv", cikti)
            self.assertEqual([(p["parti"], p["karar"], p["uygunsuz"], p["eksik"], p["tarih"]) for p in r["partiler"]],
                             [("P-101", "Uygun", 0, [], date(2026, 10, 8)), ("P-102", "Sınıra yakın", 0, [], date(2026, 10, 9)), ("P-103", "Uygunsuz", 2, [], date(2026, 10, 10)),
                              ("P-104", "Eksik analiz", 0, ["Koliform"], date(2026, 10, 11)), ("P-105", "Uygunsuz", 1, [], date(2026, 10, 12))])
            self.assertEqual([(u["tur"], u["kim"]) for u in r["uyarilar"]], [("Eksik analiz", "P-104"), ("Uygunsuz", "P-103 · Kuru Madde"), ("Uygunsuz", "P-103 · Tuz"), ("Uygunsuz", "P-105 · Koliform"), ("Sınıra yakın", "P-102 · pH")])
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Parti Özeti", "Sonuçlar", "Parametre Özeti", "Uyarılar"])
            self.assertEqual([c.value for c in wb["Sonuçlar"][5]][:7], ["P-101", "Örnek Beyaz Peynir", "Koliform", "<10", "kob/g", "≤ 100", "Uygun"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            sy, py = Path(tmp) / "s.csv", Path(tmp) / "p.csv"
            py.write_text("Ürün;Parametre;Min;Max;Beklenen\nU;pH;4,3;4,8;\nU;pH;1;2;\nU;Nem;x;;\nU;Kül;;;\nU;Yağ;5;3;\n;Küf;;;Yok\n", encoding="utf-8")
            sy.write_text("Numune No;Ürün;Analiz;Değer\nN1;U;pH;4,5\nN1;U;Küf;Yok\nN1;U;;3\nN2;U;pH;\n", encoding="utf-8")
            spec, uy = main.oku_spec(py)
            self.assertEqual((sorted(spec), sorted(u["tur"] for u in uy)), ([("", "kuf"), ("u", "ph")], ["Mükerrer spesifikasyon"] + ["Spesifikasyon okunamadı"] * 3))
            l, uy = main.oku_sonuc(sy)
            self.assertEqual(([(x["parti"], x["parametre"], x["sonuc"]) for x in l], [u["tur"] for u in uy]), ([("N1", "pH", "4,5"), ("N1", "Küf", "Yok")], ["Sonuç satırı okunamadı"] * 2))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--sonuclar", str(sy), "--spesifikasyon", str(py), "--cikti", c]
            self.assertEqual(main.main(arg + ["--yakin", "5"]), 0)
            self.assertEqual(main.main(arg + ["--yakin", "50"]), 2)
            self.assertEqual(main.main(["--sonuclar", str(sy), "--cikti", c]), 2)
            self.assertEqual(main.main(["--sonuclar", str(Path(tmp) / "yok.csv"), "--spesifikasyon", str(py)]), 1)
            py.write_text("Ürün;Parametre;Min;Max\nU;pH;x;\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            sy.write_text("Parti;Sonuç\n1;2\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
