import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def sip(*satir):
    return {("s1", main.katla(r), main.katla(b)): {"siparis": "S1", "model": "M", "renk": r, "beden": b, "adet": D(a)} for r, b, a in satir}


def kes(*satir):
    return {("s1", main.katla(r), main.katla(b)): D(a) for r, b, a in satir}


def kum(renk="Beyaz", verilen=1300, iade=40, sarfiyat="1.2"):
    return {("s1", main.katla(renk)): {"verilen": D(verilen), "iade": D(iade), "sarfiyat": None if sarfiyat is None else D(sarfiyat), "birim": "m"}}


class AnalizTesti(unittest.TestCase):
    def test_kesim_durumu(self):
        r = main.analiz(sip(("Beyaz", "S", 200), ("Beyaz", "M", 100), ("Beyaz", "L", 100), ("Beyaz", "XL", 100), ("Beyaz", "XXL", 100)),
                        kes(("Beyaz", "S", 206), ("Beyaz", "M", 99), ("Beyaz", "L", 105), ("Beyaz", "XL", 106)), {})
        self.assertEqual([(x["beden"], x["fark"], x["oran"], x["durum"]) for x in r["satirlar"]],
                         [("S", D(6), D("103.00"), "Uygun"), ("M", D(-1), D("99.00"), "Eksik kesim"), ("L", D(5), D("105.00"), "Uygun"), ("XL", D(6), D("106.00"), "Fazla kesim"),
                          ("XXL", D(-100), D("0.00"), "Eksik kesim")])
        t = r["toplam"]
        self.assertEqual((t["adet"], t["kesilen"], t["oran"], t["eksik"], t["fazla"], t["fire"]), (D(600), D(516), D("86.00"), 2, 1, None))
        r = main.analiz(sip(("Beyaz", "M", 100)), kes(("Beyaz", "M", 99)), {}, eksik_tol=D(1))
        self.assertEqual(r["satirlar"][0]["durum"], "Uygun")

    def test_kumas_ve_fire(self):
        s, k = sip(("Beyaz", "S", 500), ("Beyaz", "M", 500)), kes(("Beyaz", "S", 512), ("Beyaz", "M", 513))
        x = main.analiz(s, k, kum())["renkler"][0]
        self.assertEqual((x["kullanilan"], x["plan"], x["fire"], x["fire_oran"], x["sarfiyat"], x["isaret"]), (D(1260), D("1230.0"), D("30.0"), D("2.38"), D("1.229"), ""))   # 1.025 adet × 1,20
        x = main.analiz(s, k, kum(iade=0))["renkler"][0]
        self.assertEqual((x["fire"], x["fire_oran"], x["isaret"]), (D("70.0"), D("5.38"), "Fire yüksek"))
        self.assertEqual(main.analiz(s, k, kum(iade=0), fire_esik=D(6))["renkler"][0]["isaret"], "")
        self.assertEqual(main.analiz(s, k, kum(iade=100))["renkler"][0]["isaret"], "Plandan az kullanım")
        x = main.analiz(s, k, kum(sarfiyat=None))["renkler"][0]
        self.assertEqual((x["kullanilan"], x["fire"], x["sarfiyat"]), (D(1260), None, D("1.229")))
        r = main.analiz(s, k, kum("Siyah"))
        self.assertEqual(sorted(u["tur"] for u in r["uyarilar"]), ["Kumaş kaydı yok", "Kumaşın siparişi yok"])
        with self.assertRaises(ValueError):
            main.analiz({}, {}, {})


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "k.xlsx"
            r = main.calistir(ORNEK / "siparis.csv", ORNEK / "kesim.csv", ORNEK / "kumas.csv", cikti)
            d = {(x["renk"], x["beden"]): (x["kesilen"], x["durum"]) for x in r["satirlar"]}
            self.assertEqual((d[("Beyaz", "M")], d[("Beyaz", "XL")], d[("Siyah", "M")], d[("Siyah", "XL")]), ((D(412), "Uygun"), (D(98), "Eksik kesim"), (D(318), "Fazla kesim"), (D(104), "Uygun")))
            k = {x["renk"]: (x["kesilen"], x["kullanilan"], x["plan"], x["fire"], x["fire_oran"], x["sarfiyat"], x["isaret"]) for x in r["renkler"]}
            self.assertEqual(k["Beyaz"], (D(1025), D("1260.00"), D("1230.00"), D("30.00"), D("2.38"), D("1.229"), ""))
            self.assertEqual(k["Siyah"], (D(822), D("1040.00"), D("986.40"), D("53.60"), D("5.15"), D("1.265"), "Fire yüksek"))
            t = r["toplam"]
            self.assertEqual((t["adet"], t["kesilen"], t["oran"], t["kullanilan"], t["fire"], t["fire_oran"], t["eksik"], t["fazla"]),
                             (D(1800), D(1847), D("102.61"), D("2300.00"), D("83.60"), D("3.63"), 1, 1))
            self.assertEqual(r["uyarilar"], [])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Kesim Föyü", "Kumaş ve Fire", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            s, k, m = (Path(tmp) / x for x in ("s.csv", "k.csv", "m.csv"))
            s.write_text("Sipariş No;Renk;Beden;Sipariş Adedi\nS1;Beyaz;M;100\nS1;beyaz;m;50\nS1;Beyaz;L;çok\n", encoding="utf-8")
            k.write_text("Sipariş No;Renk;Beden;Kesilen Adet\nS1;Beyaz;M;80\nS1;BEYAZ;M;75\nS1;Beyaz;XL;10\nS1;Beyaz;M;yok\n", encoding="utf-8")
            m.write_text("Sipariş No;Renk;Verilen Kumaş;İade;Planlanan Birim Sarfiyat\nS1;Beyaz;200;10;az\nS1;Siyah;100;150;1\n", encoding="utf-8")
            sp, ks, km, uy = main.oku(s, k, m)
            a = ("s1", "beyaz", "m")
            self.assertEqual((sp[a]["adet"], ks[a], km[("s1", "beyaz")]["sarfiyat"], len(sp), len(km)), (D(150), D(155), None, 1, 1))
            self.assertEqual(sorted(u["tur"] for u in uy), ["Kesilen adet okunamadı", "Kumaş miktarı okunamadı", "Planlanan sarfiyat okunamadı", "Sipariş adedi okunamadı", "Siparişte olmayan kesim"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--siparis", str(s), "--kesim", str(k), "--cikti", c]
            self.assertEqual(main.main(arg + ["--kumas", str(m), "--fazla-tol", "3"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(arg + ["--fire-esik", "-1"]), 2)
            self.assertEqual(main.main(["--siparis", str(Path(tmp) / "yok.csv")]), 1)
            s.write_text("Sipariş No;Beden;Sipariş Adedi\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
