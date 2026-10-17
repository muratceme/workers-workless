"""
Kesim Föyü ve Fire Raporu — Gerisi Bende kod bloğu
Üretim ve Fason Takip › Kesimhane Şefi

Kesilen adetleri beden-renk bazında sipariş adetleriyle karşılaştırır; kumaş kullanımını ve kesim firesini hesaplar.
  Kesim farkı       = kesilen adet − sipariş adedi;   kesim % = kesilen ÷ sipariş × 100
  Durum             = kesim % < 100 − --eksik-tol (0) "Eksik kesim";  > 100 + --fazla-tol (5) "Fazla kesim";  arası "Uygun"
  Kullanılan kumaş  = verilen − iade
  Planlanan kullanım = kesilen adet × planlanan birim sarfiyat (pastal planı)
  Kesim firesi      = kullanılan − planlanan kullanım;   fire % = fire ÷ kullanılan × 100;   > --fire-esik (3) işaretlenir
  Gerçekleşen birim sarfiyat = kullanılan ÷ kesilen adet
Aynı sipariş-renk-beden için birden çok kesim satırı (pastal) toplanır. Kumaş birimi (metre / kg) dosyada tutarlı olmalıdır.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SIP_SUTUN = {"siparis": ("siparis no", "siparis", "order no"), "model": ("model", "model no"), "renk": ("renk",), "beden": ("beden",), "adet": ("siparis adedi", "adet", "siparis adet")}
KES_SUTUN = {"siparis": ("siparis no", "siparis", "order no"), "renk": ("renk",), "beden": ("beden",), "adet": ("kesilen adet", "kesilen", "kesim adedi", "adet"), "pastal": ("pastal no", "pastal", "kesim no")}
KUM_SUTUN = {"siparis": ("siparis no", "siparis", "order no"), "renk": ("renk",), "verilen": ("verilen kumas", "verilen", "cikan kumas", "depodan cikan"), "iade": ("iade", "iade kumas", "kalan kumas"),
             "sarfiyat": ("planlanan birim sarfiyat", "birim sarfiyat", "planlanan sarfiyat", "pastal sarfiyati"), "birim": ("birim",)}
Z = Decimal(0)


def oku(sip_yol: Path, kes_yol: Path, kum_yol: Path | None) -> tuple[dict, dict, dict, list[dict]]:
    uy, sip, kes, kum = [], {}, defaultdict(lambda: Z), {}
    for r in kayitlar(sip_yol, SIP_SUTUN, ("siparis", "beden", "adet")):
        if not metin(r["siparis"]):
            continue
        a = para(r["adet"])
        k = (katla(r["siparis"]), katla(r["renk"]), katla(r["beden"]))
        if a is None or a < 0:
            uy.append(uyari("Yüksek", "Sipariş adedi okunamadı", f"{metin(r['siparis'])} · {metin(r['renk'])} · {metin(r['beden'])}", f"{sip_yol.name} satır {r['_satir']}: '{metin(r['adet'])}'."))
            continue
        if k in sip:
            sip[k]["adet"] += a
        else:
            sip[k] = {"siparis": metin(r["siparis"]), "model": metin(r["model"]), "renk": metin(r["renk"]) or "—", "beden": metin(r["beden"]), "adet": a}
    pastal = defaultdict(set)
    for r in kayitlar(kes_yol, KES_SUTUN, ("siparis", "beden", "adet")):
        if not metin(r["siparis"]):
            continue
        a = para(r["adet"])
        k = (katla(r["siparis"]), katla(r["renk"]), katla(r["beden"]))
        if a is None or a < 0:
            uy.append(uyari("Yüksek", "Kesilen adet okunamadı", f"{metin(r['siparis'])} · {metin(r['renk'])} · {metin(r['beden'])}", f"{kes_yol.name} satır {r['_satir']}: '{metin(r['adet'])}'."))
            continue
        if k not in sip:
            uy.append(uyari("Yüksek", "Siparişte olmayan kesim", f"{metin(r['siparis'])} · {metin(r['renk'])} · {metin(r['beden'])}", f"{kes_yol.name} satır {r['_satir']}: {a:f} adet kesim, sipariş listesinde karşılığı yok; rapora alınmadı."))
            continue
        kes[k] += a
        pastal[k].add(metin(r["pastal"]))
    if kum_yol:
        for r in kayitlar(kum_yol, KUM_SUTUN, ("siparis", "verilen")):
            if not metin(r["siparis"]):
                continue
            kim = f"{metin(r['siparis'])} · {metin(r['renk'])}"
            v, i, s = para(r["verilen"]), para(r["iade"]) or Z, para(r["sarfiyat"])
            if v is None or v <= 0 or i < 0 or i > v:
                uy.append(uyari("Yüksek", "Kumaş miktarı okunamadı", kim, f"{kum_yol.name} satır {r['_satir']}: verilen '{metin(r['verilen'])}', iade '{metin(r['iade'])}'. İade verilenden büyük olamaz."))
                continue
            if s is None and metin(r["sarfiyat"]):
                uy.append(uyari("Orta", "Planlanan sarfiyat okunamadı", kim, f"'{metin(r['sarfiyat'])}'. Fire hesaplanmadı."))
            kum[(katla(r["siparis"]), katla(r["renk"]))] = {"verilen": v, "iade": i, "sarfiyat": s if s and s > 0 else None, "birim": metin(r["birim"])}
    return sip, dict(kes) | {k: Z for k in sip if k not in kes}, kum, uy


def analiz(sip: dict, kes: dict, kum: dict, eksik_tol: Decimal = Z, fazla_tol: Decimal = Decimal(5), fire_esik: Decimal = Decimal(3)) -> dict:
    if not sip:
        raise ValueError("Sipariş satırı yok.")
    uy, satirlar = [], []
    for k, s in sip.items():
        c = kes.get(k, Z)
        oran = yuvarla(c / s["adet"] * 100) if s["adet"] else None
        durum = "Sipariş 0" if oran is None else "Eksik kesim" if oran < 100 - eksik_tol else "Fazla kesim" if oran > 100 + fazla_tol else "Uygun"
        satirlar.append(s | {"kesilen": c, "fark": c - s["adet"], "oran": oran, "durum": durum})
    renkler = {}
    for x in satirlar:
        r = renkler.setdefault((katla(x["siparis"]), katla(x["renk"])), {"siparis": x["siparis"], "model": x["model"], "renk": x["renk"], "adet": Z, "kesilen": Z, "eksik": 0, "fazla": 0})
        r["adet"] += x["adet"]
        r["kesilen"] += x["kesilen"]
        r["eksik"] += x["durum"] == "Eksik kesim"
        r["fazla"] += x["durum"] == "Fazla kesim"
    for k, r in renkler.items():
        r["oran"] = yuvarla(r["kesilen"] / r["adet"] * 100) if r["adet"] else None
        m = kum.get(k)
        r |= {"verilen": None, "iade": None, "kullanilan": None, "plan_sarfiyat": None, "plan": None, "fire": None, "fire_oran": None, "sarfiyat": None, "birim": "", "isaret": ""}
        if not m:
            if kum:
                uy.append(uyari("Orta", "Kumaş kaydı yok", f"{r['siparis']} · {r['renk']}", "Bu sipariş ve renk için kumaş satırı bulunamadı; kumaş kullanımı hesaplanmadı."))
            continue
        r |= {"verilen": m["verilen"], "iade": m["iade"], "kullanilan": m["verilen"] - m["iade"], "plan_sarfiyat": m["sarfiyat"], "birim": m["birim"]}
        if r["kesilen"] > 0:
            r["sarfiyat"] = yuvarla(r["kullanilan"] / r["kesilen"], "0.001")
            if m["sarfiyat"]:
                r["plan"] = r["kesilen"] * m["sarfiyat"]
                r["fire"] = r["kullanilan"] - r["plan"]
                r["fire_oran"] = yuvarla(r["fire"] / r["kullanilan"] * 100) if r["kullanilan"] else None
                r["isaret"] = "Fire yüksek" if r["fire_oran"] is not None and r["fire_oran"] > fire_esik else "Plandan az kullanım" if r["fire"] < 0 else ""
    for k in kum:
        if k not in renkler:
            uy.append(uyari("Orta", "Kumaşın siparişi yok", f"{k[0]} · {k[1]}", "Kumaş dosyasında var, sipariş listesinde bu sipariş ve renk yok."))
    rl = list(renkler.values())
    kl = [r for r in rl if r["kullanilan"] is not None]
    fl = [r for r in rl if r["fire"] is not None]
    t = {"adet": sum((r["adet"] for r in rl), Z), "kesilen": sum((r["kesilen"] for r in rl), Z), "kullanilan": sum((r["kullanilan"] for r in kl), Z) if kl else None,
         "plan": sum((r["plan"] for r in fl), Z) if fl else None, "fire": sum((r["fire"] for r in fl), Z) if fl else None,
         "eksik": sum(x["durum"] == "Eksik kesim" for x in satirlar), "fazla": sum(x["durum"] == "Fazla kesim" for x in satirlar)}
    t["oran"] = yuvarla(t["kesilen"] / t["adet"] * 100) if t["adet"] else None
    fk = sum((r["kullanilan"] for r in fl), Z)
    t["fire_oran"] = yuvarla(t["fire"] / fk * 100) if fl and fk else None
    return {"satirlar": satirlar, "renkler": rl, "toplam": t, "uyarilar": uy, "esik": (eksik_tol, fazla_tol, fire_esik)}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb, t = Workbook(), r["toplam"]
    renk = {"Eksik kesim": KIRMIZI, "Fazla kesim": SARI, "Uygun": None, "Sipariş 0": None}
    ws = wb.active
    ws.title = "Kesim Föyü"
    _baslik(ws, ["Sipariş No", "Model", "Renk", "Beden", "Sipariş Adedi", "Kesilen Adet", "Fark", "Kesim %", "Durum"], (14, 14, 14, 9, 12, 12, 9, 10, 14))
    for x in r["satirlar"]:
        _satir(ws, [x["siparis"], x["model"], x["renk"], x["beden"], x["adet"], x["kesilen"], x["fark"], x["oran"], x["durum"]], para_s=(8,), dolgu={9: renk[x["durum"]]})
    _satir(ws, ["Toplam", "", "", "", t["adet"], t["kesilen"], t["kesilen"] - t["adet"], t["oran"], f"{t['eksik']} eksik, {t['fazla']} fazla"], para_s=(8,))
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.append([])
    ws.append([f"Eşikler: eksik toleransı %{r['esik'][0]}, fazla toleransı %{r['esik'][1]}, fire eşiği %{r['esik'][2]}."])
    ws = wb.create_sheet("Kumaş ve Fire")
    _baslik(ws, ["Sipariş No", "Model", "Renk", "Sipariş Adedi", "Kesilen Adet", "Kesim %", "Verilen Kumaş", "İade", "Kullanılan", "Planlanan Birim Sarfiyat", "Planlanan Kullanım",
                 "Kesim Firesi", "Fire %", "Gerçekleşen Birim Sarfiyat", "Birim", "İşaret", "Neden (siz doldurun)"], (14, 14, 14, 11, 11, 9, 13, 10, 13, 13, 13, 12, 9, 13, 8, 20, 30))
    for x in r["renkler"]:
        _satir(ws, [x["siparis"], x["model"], x["renk"], x["adet"], x["kesilen"], x["oran"], x["verilen"], x["iade"], x["kullanilan"], x["plan_sarfiyat"], x["plan"], x["fire"], x["fire_oran"],
                    x["sarfiyat"], x["birim"], x["isaret"], ""], para_s=(6, 7, 8, 9, 11, 12, 13), dolgu={16: KIRMIZI if x["isaret"] == "Fire yüksek" else SARI if x["isaret"] else None})
    _satir(ws, ["Toplam", "", "", t["adet"], t["kesilen"], t["oran"], None, None, t["kullanilan"], None, t["plan"], t["fire"], t["fire_oran"]], para_s=(6, 9, 11, 12, 13))
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(sip_yol: Path, kes_yol: Path, kum_yol: Path | None, cikti: Path, **esik) -> dict:
    sip, kes, kum, uy = oku(sip_yol, kes_yol, kum_yol)
    r = analiz(sip, kes, kum, **esik)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, r["uyarilar"])
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Kesilen adetleri siparişle karşılaştırır; kumaş kullanımını ve kesim firesini hesaplar.")
    p.add_argument("--siparis", type=Path, default=ORNEK / "siparis.csv", help="Sipariş No, Model, Renk, Beden, Sipariş Adedi")
    p.add_argument("--kesim", type=Path, default=ORNEK / "kesim.csv", help="Sipariş No, Renk, Beden, Kesilen Adet, Pastal No")
    p.add_argument("--kumas", type=Path, help="Sipariş No, Renk, Verilen Kumaş, İade, Planlanan Birim Sarfiyat, Birim")
    p.add_argument("--eksik-tol", default="0", help="Kabul edilen eksik kesim %% (varsayılan 0)")
    p.add_argument("--fazla-tol", default="5", help="Kabul edilen fazla kesim %% (varsayılan 5)")
    p.add_argument("--fire-esik", default="3", help="Bu kesim firesi %%'sinin üstü işaretlenir (varsayılan 3)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kesim_foyu.xlsx")
    a = p.parse_args(argv)
    kumas = a.kumas or (ORNEK / "kumas.csv" if a.siparis == ORNEK / "siparis.csv" else None)
    for y in (a.siparis, a.kesim, kumas):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    e = [para(x) for x in (a.eksik_tol, a.fazla_tol, a.fire_esik)]
    if any(x is None or x < 0 for x in e):
        print("[X] --eksik-tol, --fazla-tol ve --fire-esik sıfır veya pozitif sayı olmalı")
        return 2
    try:
        r = calistir(a.siparis, a.kesim, kumas, a.cikti, eksik_tol=e[0], fazla_tol=e[1], fire_esik=e[2])
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] Sipariş {t['adet']:f} · kesilen {t['kesilen']:f} (%{t['oran']}) · {t['eksik']} eksik, {t['fazla']} fazla satır" + (f" · kesim firesi %{t['fire_oran']}" if t["fire_oran"] is not None else ""))
    for x in r["satirlar"]:
        if x["durum"] == "Eksik kesim":
            print(f"[!] Eksik kesim · {x['siparis']} {x['renk']} {x['beden']}: {x['kesilen']:f} / {x['adet']:f}")
    for x in r["renkler"]:
        if x["isaret"] == "Fire yüksek":
            print(f"[!] Fire yüksek · {x['siparis']} {x['renk']}: %{x['fire_oran']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
