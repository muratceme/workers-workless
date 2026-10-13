import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def b(bayi, satis, hedef=None, tahsilat=None, alacak=None, stok=None, cikis=None, onceki=None, bolge="X"):
    f = lambda v: None if v is None else D(v)   # noqa: E731
    return {"bayi": bayi, "bolge": bolge, "satis": D(satis), "hedef": f(hedef), "tahsilat": f(tahsilat), "alacak": f(alacak), "stok": f(stok), "cikis": f(cikis), "onceki": f(onceki)}


class AnalizTesti(unittest.TestCase):
    def test_gostergeler(self):
        x = main.analiz([b("A", 1200, 1000, 1050, 900, 600, 1100, 1000), b("B", 300)])["bayiler"]
        self.assertEqual([(v["bayi"], v["gerceklesme"], v["buyume"], v["tahsilat_y"], v["alacak_g"], v["stok_g"], v["pay"], v["isaret"]) for v in x],
                         [("A", D("120.00"), D("20.00"), D("87.50"), D("22.50"), D("16.36"), D("80.00"), []), ("B", None, None, None, None, None, D("20.00"), [])])
        x = main.analiz([b("A", 450, stok=200, alacak=450)], gun=90)["bayiler"][0]
        self.assertEqual((x["stok_g"], x["alacak_g"]), (D("40.00"), D("90.00")))                           # bayi satışı yoksa bayiye satışa göre

    def test_isaretler_ve_esikler(self):
        AY = ("gun", "hedef_esik", "alacak_gun", "stok_gun")
        i = lambda *a, **kw: main.analiz([b("A", *a, **{k: v for k, v in kw.items() if k not in AY})], **{k: v for k, v in kw.items() if k in AY})["bayiler"][0]["isaret"]   # noqa: E731
        self.assertEqual((i(799, 1000), i(800, 1000), i(100, alacak=200), i(100, alacak=201), i(100, stok=150), i(100, stok=151), i(100, stok=151, cikis=200)),
                         (["Hedefin altında"], [], [], ["Alacak günü yüksek"], [], ["Stok günü yüksek"], []))
        self.assertEqual((i(0, 100, alacak=5), i(0, 100), i(0), i(0, stok=5, cikis=1), i(90, 100, hedef_esik=D(95)), i(100, alacak=100, alacak_gun=D(29)), i(100, stok=100, stok_gun=D(29))),
                         (["Hedefin altında", "Satış yok"], ["Hedefin altında"], [], ["Stok günü yüksek", "Satış yok"], ["Hedefin altında"], ["Alacak günü yüksek"], ["Stok günü yüksek"]))
        r = main.analiz([b("A", 0, 100, alacak=5)])
        self.assertEqual(sorted(u["tur"] for u in r["uyarilar"]), ["Hedefin altında", "Satış yok, alacak veya stok var"])

    def test_bolge_ve_toplam(self):
        r = main.analiz([b("A", 100, 200, 80, 50, bolge="Ege"), b("B", 300, 200, 300, 400, 100, 200, 250, bolge="EGE"), b("C", 600, 500, bolge="Marmara"), b("D", 0, bolge="")])
        self.assertEqual([x["bayi"] for x in r["bayiler"]], ["C", "B", "A", "D"])
        self.assertEqual([(o["bolge"], o["bayi"], o["isaretli"], o["satis"], o["hedef"], o["gerceklesme"], o["tahsilat_y"], o["alacak_g"], o["stok_g"], o["buyume"], o["pay"]) for o in r["bolgeler"]],
                         [("Marmara", 1, 0, D(600), D(500), D("120.00"), None, None, None, None, D("60.00")), ("Ege", 2, 1, D(400), D(400), D("100.00"), D("95.00"), D("33.75"), D("15.00"), D("20.00"), D("40.00")),
                          ("", 1, 0, D(0), None, None, None, None, None, None, D("0.00"))])
        t = r["toplam"]
        self.assertEqual((t["bayi"], t["isaretli"], t["satis"], t["hedef"], t["gerceklesme"], t["tahsilat"], t["alacak"], t["alacak_g"], t["pay"]), (4, 1, D(1000), D(900), D("111.11"), D(380), D(450), D("13.50"), D("100.00")))
        with self.assertRaises(ValueError):
            main.analiz([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "b.xlsx"
            r = main.calistir(ORNEK / "bayiler.csv", cikti)
            self.assertEqual([(x["bayi"], x["gerceklesme"], x["buyume"], x["tahsilat_y"], x["alacak_g"], x["stok_g"], x["isaret"]) for x in r["bayiler"]],
                             [("Örnek Dağıtım A.Ş.", D("120.00"), D("20.00"), D("87.50"), D("22.50"), D("16.36"), []), ("Temsili Toptan A.Ş.", D("105.56"), None, D("94.74"), D("22.11"), D("15.00"), []),
                              ("Deneme Ticaret Ltd.", D("80.00"), D("-11.11"), D("62.50"), D("71.25"), D("64.29"), ["Alacak günü yüksek", "Stok günü yüksek"]),
                              ("Numune Pazarlama Ltd.", D("75.00"), D("12.50"), D("100.00"), D("20.00"), D("13.33"), ["Hedefin altında"]),
                              ("Misal Gıda Dağıtım Ltd.", D("0.00"), D("-100.00"), None, None, D("90.00"), ["Hedefin altında", "Stok günü yüksek", "Satış yok"])])
            t = r["toplam"]
            self.assertEqual((t["satis"], t["hedef"], t["gerceklesme"], t["isaretli"], [o["bolge"] for o in r["bolgeler"]]), (D("3400000.00"), D("3800000.00"), D("89.47"), 3, ["Marmara", "İç Anadolu", "Ege"]))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Bayiler", "Bölge Özeti", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "b.csv"
            y.write_text("Distribütör;İl;Net Satış;Hedef;Alacak\nA;İzmir;1.250,50;;\nA;İzmir;100;200;x\nB;;;100;\nC;;-5;;\nA;İzmir;10;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["bayi"], x["bolge"], x["satis"], x["hedef"], x["stok"]) for x in l], sorted(u["tur"] for u in uy)),
                             ([("A", "İzmir", D("1250.50"), None, None), ("A", "İzmir", D(10), None, None)], ["Bayi satırı okunamadı"] * 3 + ["Mükerrer bayi"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--bayiler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--gun", "90", "--hedef-esik", "70", "--alacak-gun", "45", "--stok-gun", "30"]), 0)
            for ek in (["--gun", "0"], ["--hedef-esik", "x"], ["--alacak-gun", "-1"]):
                self.assertEqual(main.main(arg + ek), 2)
            self.assertEqual(main.main(["--bayiler", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Bayi;Hedef\nA;1\n", "Bayi;Satış\nA;x\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
