"""
Bayi ve Distribütör Satış Raporu — Workers / Workless kod bloğu
Satış ve Kanal Yönetimi › Kanal Satış Yöneticisi

Bayi bazında satış, hedef, tahsilat, alacak ve stok verisinden hedef gerçekleşme, tahsilat oranı, alacak ve stok günlerini raporlar.
  Hedef gerçekleşme % = satış ÷ hedef × 100
  Büyüme %            = (satış − önceki dönem satışı) ÷ önceki dönem satışı × 100   (bölge ve toplamda yalnız önceki dönemi olan bayiler karşılaştırılır)
  Tahsilat oranı %    = dönem tahsilatı ÷ dönem satışı × 100                    (önceki dönem alacağı tahsil edildiyse %100'ü aşabilir)
  Alacak gün          = açık alacak ÷ satış × dönem gün sayısı                  (--gun, varsayılan 30)
  Stok gün            = bayi stoku ÷ bayinin kendi satışı × dönem gün sayısı    (bayi satışı yoksa bayiye yapılan satış kullanılır)
  Satış payı %        = bayi satışı ÷ toplam satış × 100
  İşaretler           = hedef gerçekleşme < --hedef-esik (%80) · alacak gün > --alacak-gun (60) · stok gün > --stok-gun (45) · satış yok ama alacak veya stok var
Tutarlar aynı esasta olmalıdır (hepsi KDV dahil veya hepsi hariç; stok ve bayi satışı aynı fiyat düzeyinde). Eşikler sizin kanal politikanıza göre parametreyle değişir.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"bayi": ("bayi", "distributor", "bayi adi", "cari", "musteri"), "bolge": ("bolge", "satis bolgesi", "il"), "satis": ("satis", "donem satisi", "net satis", "bayiye satis", "sell in"),
         "hedef": ("hedef", "satis hedefi", "donem hedefi"), "tahsilat": ("tahsilat", "donem tahsilati"), "alacak": ("acik alacak", "alacak", "acik bakiye", "bakiye"),
         "stok": ("bayi stoku", "stok", "stok tutari"), "cikis": ("bayi satisi", "sell out", "bayi cikisi", "bayinin satisi"), "onceki": ("onceki donem satis", "gecen donem satis", "gecen yil satis", "onceki donem")}
Z = Decimal(0)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("bayi", "satis")):
        b = metin(r["bayi"])
        if not b:
            continue
        d = {a: para(r[a]) if metin(r[a]) else None for a in ("satis", "hedef", "tahsilat", "alacak", "stok", "cikis", "onceki")}
        bozuk = [a for a in d if (d[a] is None and metin(r[a])) or (d[a] is not None and d[a] < 0)]
        if d["satis"] is None or bozuk:
            uy.append(uyari("Yüksek", "Bayi satırı okunamadı", f"{b} (satır {r['_satir']})", f"Okunamayan veya negatif alan: {', '.join(bozuk) or 'satis'}. Satır alınmadı; toplamlar bu bayi olmadan hesaplandı."))
            continue
        if katla(b) in gorulen:
            uy.append(uyari("Orta", "Mükerrer bayi", b, f"Bayi birden fazla satırda (satır {r['_satir']}). Satırlar ayrı ayrı raporlandı; toplamlar ikisini de içerir."))
        gorulen.add(katla(b))
        l.append({"bayi": b, "bolge": metin(r["bolge"]) or "Belirtilmemiş"} | d)
    return l, uy


def oran(a: Decimal | None, b: Decimal | None, carpan: Decimal = Decimal(100)) -> Decimal | None:
    return yuvarla(a / b * carpan) if a is not None and b else None


def analiz(bayiler: list[dict], gun: int = 30, hedef_esik: Decimal = Decimal(80), alacak_gun: Decimal = Decimal(60), stok_gun: Decimal = Decimal(45)) -> dict:
    if not bayiler:
        raise ValueError("Raporlanacak bayi yok.")
    uy, liste, G = [], [], Decimal(gun)
    top_satis = sum((b["satis"] for b in bayiler), Z)

    def gostergeler(x: dict) -> dict:
        return {"gerceklesme": oran(x["satis"], x["hedef"]), "buyume": oran(x.get("satis_k", x["satis"]) - x["onceki"], x["onceki"]) if x["onceki"] else None, "tahsilat_y": oran(x["tahsilat"], x["satis"]),
                "alacak_g": oran(x["alacak"], x["satis"], G), "stok_g": oran(x["stok"], x["cikis"] or x["satis"], G), "pay": oran(x["satis"], top_satis)}

    bolge = {}
    for b in bayiler:
        x = b | gostergeler(b)
        isaret = []
        if x["gerceklesme"] is not None and x["gerceklesme"] < hedef_esik:
            isaret.append("Hedefin altında")
            uy.append(uyari("Orta", "Hedefin altında", b["bayi"], f"Satış {b['satis']:,.2f}, hedef {b['hedef']:,.2f} (gerçekleşme %{x['gerceklesme']})."))
        if x["alacak_g"] is not None and x["alacak_g"] > alacak_gun:
            isaret.append("Alacak günü yüksek")
            uy.append(uyari("Yüksek", "Alacak günü yüksek", b["bayi"], f"Açık alacak {b['alacak']:,.2f}, dönem satışı {b['satis']:,.2f}: {x['alacak_g']} gün (eşik {alacak_gun.normalize():f}). Sevkiyat ve limit kararını gözden geçirin."))
        if x["stok_g"] is not None and x["stok_g"] > stok_gun:
            isaret.append("Stok günü yüksek")
            uy.append(uyari("Orta", "Stok günü yüksek", b["bayi"], f"Bayi stoku {b['stok']:,.2f}: {x['stok_g']} gün (eşik {stok_gun.normalize():f}). Bayiye yeni satış stoku şişirebilir."))
        if not b["satis"] and (b["alacak"] or b["stok"]):
            isaret.append("Satış yok")
            uy.append(uyari("Yüksek", "Satış yok, alacak veya stok var", b["bayi"], f"Dönem satışı sıfır; açık alacak {b['alacak'] or Z:,.2f}, bayi stoku {b['stok'] or Z:,.2f}."))
        liste.append(x | {"isaret": isaret})
        o = bolge.setdefault(katla(b["bolge"]), {"bolge": b["bolge"], "bayi": 0, "isaretli": 0, "satis_k": Z} | {a: None for a in ("satis", "hedef", "tahsilat", "alacak", "stok", "cikis", "onceki")})
        o["bayi"] += 1
        o["isaretli"] += bool(isaret)
        o["satis_k"] += b["satis"] if b["onceki"] is not None else Z                              # büyüme yalnız önceki dönemi olan bayiler üzerinden
        for a in ("satis", "hedef", "tahsilat", "alacak", "stok", "cikis", "onceki"):
            if b[a] is not None:
                o[a] = (o[a] or Z) + b[a]
    liste.sort(key=lambda x: (-x["satis"], x["bayi"]))
    bolgeler = sorted((o | gostergeler(o) for o in bolge.values()), key=lambda o: (-o["satis"], o["bolge"]))
    toplam = {"bolge": "Toplam", "bayi": len(liste), "isaretli": sum(1 for x in liste if x["isaret"]), "satis_k": sum((b["satis"] for b in bayiler if b["onceki"] is not None), Z)}
    for a in ("satis", "hedef", "tahsilat", "alacak", "stok", "cikis", "onceki"):
        d = [b[a] for b in bayiler if b[a] is not None]
        toplam[a] = sum(d, Z) if d else None
    return {"bayiler": liste, "bolgeler": bolgeler, "toplam": toplam | gostergeler(toplam), "ayar": (gun, hedef_esik, alacak_gun, stok_gun), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    gun, he, ag, sg = r["ayar"]
    ws = wb.active
    ws.title = "Bayiler"
    _baslik(ws, ["Bayi", "Bölge", "Satış", "Hedef", "Gerçekleşme %", "Önceki Dönem", "Büyüme %", "Satış Payı %", "Tahsilat", "Tahsilat Oranı %", "Açık Alacak", "Alacak Gün", "Bayi Stoku", "Bayi Satışı", "Stok Gün", "İşaretler", "Aksiyon (siz doldurun)"],
            (30, 16, 15, 15, 11, 15, 9, 9, 15, 10, 15, 9, 15, 15, 9, 34, 30))
    for x in r["bayiler"]:
        _satir(ws, [x["bayi"], x["bolge"], x["satis"], x["hedef"], x["gerceklesme"], x["onceki"], x["buyume"], x["pay"], x["tahsilat"], x["tahsilat_y"], x["alacak"], x["alacak_g"], x["stok"], x["cikis"], x["stok_g"], ", ".join(x["isaret"]), ""],
               para_s=tuple(range(3, 16)), sar=(16,),
               dolgu={5: KIRMIZI if "Hedefin altında" in x["isaret"] else YESIL if x["gerceklesme"] is not None and x["gerceklesme"] >= 100 else None, 12: KIRMIZI if "Alacak günü yüksek" in x["isaret"] else None,
                      15: SARI if "Stok günü yüksek" in x["isaret"] else None, 3: KIRMIZI if "Satış yok" in x["isaret"] else None})
    ws.append([])
    ws.append([f"Dönem {gun} gün. İşaret eşikleri: gerçekleşme < %{he.normalize():f}, alacak gün > {ag.normalize():f}, stok gün > {sg.normalize():f}. Stok gün, bayi satışı verilmişse ona, verilmemişse bayiye yapılan satışa göre hesaplanır."])
    ws = wb.create_sheet("Bölge Özeti")
    _baslik(ws, ["Bölge", "Bayi Sayısı", "İşaretli Bayi", "Satış", "Hedef", "Gerçekleşme %", "Büyüme %", "Satış Payı %", "Tahsilat", "Tahsilat Oranı %", "Açık Alacak", "Alacak Gün", "Bayi Stoku", "Stok Gün"],
            (22, 9, 9, 15, 15, 11, 9, 9, 15, 10, 15, 9, 15, 9))
    for o in r["bolgeler"] + [r["toplam"]]:
        n = _satir(ws, [o["bolge"], o["bayi"], o["isaretli"] or None, o["satis"], o["hedef"], o["gerceklesme"], o["buyume"], o["pay"], o["tahsilat"], o["tahsilat_y"], o["alacak"], o["alacak_g"], o["stok"], o["stok_g"]], para_s=tuple(range(4, 15)))
        if o is r["toplam"]:
            for c in ws[n]:
                c.font = Font(bold=True)
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, **ayar) -> dict:
    bayiler, uy = oku(yol)
    r = analiz(bayiler, **ayar)
    uy += r["uyarilar"]
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Bayi bazında satış, hedef, tahsilat, alacak ve stok günlerini raporlar.")
    p.add_argument("--bayiler", type=Path, default=ORNEK / "bayiler.csv", help="Bayi, Bölge, Satış, Hedef, Tahsilat, Açık Alacak, Bayi Stoku, Bayi Satışı, Önceki Dönem Satış")
    p.add_argument("--gun", type=int, default=30, help="Dönem gün sayısı (varsayılan 30)")
    p.add_argument("--hedef-esik", default="80", help="Gerçekleşme bu %%'nin altındaysa işaretlenir (varsayılan 80)")
    p.add_argument("--alacak-gun", default="60", help="Alacak gün eşiği (varsayılan 60)")
    p.add_argument("--stok-gun", default="45", help="Stok gün eşiği (varsayılan 45)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "bayi_satis_raporu.xlsx")
    a = p.parse_args(argv)
    he, ag, sg = para(a.hedef_esik), para(a.alacak_gun), para(a.stok_gun)
    if a.gun <= 0 or any(v is None or v < 0 for v in (he, ag, sg)):
        print("[X] --gun pozitif tam sayı; --hedef-esik, --alacak-gun ve --stok-gun sıfır veya pozitif sayı olmalı")
        return 2
    if not a.bayiler.exists():
        print(f"[X] Dosya bulunamadı: {a.bayiler}")
        return 1
    try:
        r = calistir(a.bayiler, a.cikti, gun=a.gun, hedef_esik=he, alacak_gun=ag, stok_gun=sg)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {t['bayi']} bayi · satış {t['satis']:,.2f}" + (f" · hedef gerçekleşme %{t['gerceklesme']}" if t["gerceklesme"] is not None else "") + f" · {t['isaretli']} bayi işaretli")
    for x in r["bayiler"]:
        if x["isaret"]:
            print(f"[!] {x['bayi']}: {', '.join(x['isaret'])}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
