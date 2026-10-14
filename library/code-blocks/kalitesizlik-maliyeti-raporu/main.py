"""
Kalitesizlik Maliyeti Raporu — Workers / Workless kod bloğu
Kalite Güvence ve Kalite Kontrol › Kalite Müdürü

Hurda, yeniden işleme, iade ve garanti maliyetlerini ürün ve neden bazında toplar.
  İç hata maliyeti     = ürün müşteriye ulaşmadan yakalanan hatalar: hurda, yeniden işleme, ayıklama, yeniden kontrol
  Dış hata maliyeti    = müşteriye ulaştıktan sonra ortaya çıkan hatalar: iade, garanti, müşteri şikâyeti, ceza
  Pay %                = kalemin tutarı ÷ toplam kalitesizlik maliyeti × 100
  Kümülatif % ve A sınıfı = büyükten küçüğe sıralı payların birikimi; kendisinden önceki kümülatif pay %80'in altındaysa A
  Dış hata payı %      = dış hata maliyeti ÷ toplam × 100;   --dis (%50) üstündeyse uyarılır
  Ciroya oran %        = toplam kalitesizlik maliyeti ÷ --ciro × 100   (ciro verilirse)
Önleme ve ölçme-değerlendirme maliyetleri bu raporda yoktur; tutarları siz verirsiniz, kod maliyet hesaplamaz.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"tarih": ("tarih", "kayit tarihi"), "urun": ("urun", "urun adi", "stok kodu", "mamul"), "tur": ("maliyet turu", "tur", "hata maliyeti turu", "kategori"), "neden": ("neden", "hata nedeni", "hata turu"),
         "adet": ("adet", "miktar"), "tutar": ("tutar", "maliyet", "maliyet tutari")}
IC, DIS, DIGER = "İç hata", "Dış hata", "Sınıflanmamış"
TUR = {"hurda": ("Hurda", IC), "fire": ("Hurda", IC), "iskarta": ("Hurda", IC), "yeniden isleme": ("Yeniden işleme", IC), "tamir": ("Yeniden işleme", IC), "rework": ("Yeniden işleme", IC), "ayiklama": ("Ayıklama", IC),
       "yeniden kontrol": ("Yeniden kontrol", IC), "iade": ("İade", DIS), "musteri iadesi": ("İade", DIS), "garanti": ("Garanti", DIS), "musteri sikayeti": ("Müşteri şikâyeti", DIS), "sikayet": ("Müşteri şikâyeti", DIS),
       "ceza": ("Ceza", DIS), "musteri cezasi": ("Ceza", DIS)}
YOK = "(neden yazılmamış)"
SIFIR = Decimal(0)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("tur", "tutar")):
        if not metin(r["tur"]) and not metin(r["tutar"]):
            continue
        tutar, adet = para(r["tutar"]), para(r["adet"]) if metin(r["adet"]) else None
        if tutar is None or tutar < 0 or not metin(r["tur"]):
            uy.append(uyari("Yüksek", "Kayıt okunamadı", f"Satır {r['_satir']}", f"Maliyet türü '{metin(r['tur'])}' ya da tutar '{metin(r['tutar'])}' okunamadı. Satır alınmadı."))
            continue
        ad, sinif = TUR.get(katla(r["tur"]), (metin(r["tur"]), DIGER))
        l.append({"tarih": tarih(r["tarih"]), "urun": metin(r["urun"]) or "(ürün yok)", "tur": ad, "sinif": sinif, "neden": metin(r["neden"]) or YOK, "adet": adet, "tutar": tutar})
    return l, uy


def _oran(pay: Decimal, payda: Decimal) -> Decimal | None:
    return yuvarla(pay * 100 / payda) if payda else None


def raporla(kayit: list[dict], ciro: Decimal | None = None, dis_esik: Decimal = Decimal(50)) -> dict:
    if not kayit:
        raise ValueError("Raporlanacak kalitesizlik maliyeti kaydı yok.")
    toplam = sum((k["tutar"] for k in kayit), SIFIR)
    if toplam == 0:
        raise ValueError("Toplam kalitesizlik maliyeti sıfır.")

    def pareto(alan: str) -> list[dict]:
        grup = {}
        for k in kayit:
            g = grup.setdefault(katla(k[alan]), {"ad": k[alan], "kayit": 0, "tutar": SIFIR, "ic": SIFIR, "dis": SIFIR, "sinif_ad": k["sinif"]})
            g["kayit"], g["tutar"] = g["kayit"] + 1, g["tutar"] + k["tutar"]
            g["ic"], g["dis"] = g["ic"] + (k["tutar"] if k["sinif"] == IC else 0), g["dis"] + (k["tutar"] if k["sinif"] == DIS else 0)
        s, kum = sorted(grup.values(), key=lambda g: (-g["tutar"], katla(g["ad"]))), SIFIR
        for g in s:
            g["pay"], g["sinif"] = _oran(g["tutar"], toplam), "A" if _oran(kum, toplam) < 80 else "B"
            kum += g["tutar"]
            g["kum"] = _oran(kum, toplam)
        return s

    turler, urunler, nedenler, uy = pareto("tur"), pareto("urun"), pareto("neden"), []
    ic, dis = sum((k["tutar"] for k in kayit if k["sinif"] == IC), SIFIR), sum((k["tutar"] for k in kayit if k["sinif"] == DIS), SIFIR)
    ozet = {"kayit": len(kayit), "toplam": toplam, "ic": ic, "dis": dis, "diger": toplam - ic - dis, "ic_pay": _oran(ic, toplam), "dis_pay": _oran(dis, toplam), "ciro": ciro, "ciro_oran": _oran(toplam, ciro) if ciro else None}
    if ozet["dis_pay"] > dis_esik:
        uy.append(uyari("Orta", "Dış hata payı yüksek", "Toplam", f"Kalitesizlik maliyetinin %{ozet['dis_pay']} kadarı müşteriye ulaşan hatalardan ({dis:,.2f}; eşik %{dis_esik}). Hatalar içeride yakalanamıyor olabilir."))
    for n in nedenler:
        if n["ad"] == YOK:
            uy.append(uyari("Orta", "Nedeni yazılmamış maliyet", YOK, f"{n['kayit']} kayıt, {n['tutar']:,.2f} (%{n['pay']}). Neden yazılmadan iyileştirme önceliği belirlenemez."))
    for t in turler:
        if t["sinif_ad"] == DIGER:
            uy.append(uyari("Bilgi", "Maliyet türü sınıflanamadı", t["ad"], f"{t['kayit']} kayıt, {t['tutar']:,.2f}. İç ya da dış hata olarak tanınmadı; '{DIGER}' altında gösterildi."))
    matris = [{"urun": u["ad"], "hucre": [sum((k["tutar"] for k in kayit if katla(k["urun"]) == katla(u["ad"]) and k["tur"] == t["ad"]), SIFIR) for t in turler], "toplam": u["tutar"]} for u in urunler]
    ay = {}
    for k in kayit:
        if k["tarih"]:
            a = ay.setdefault(f"{k['tarih']:%Y-%m}", {"ic": SIFIR, "dis": SIFIR, "toplam": SIFIR})
            a["toplam"] += k["tutar"]
            if k["sinif"] in (IC, DIS):
                a["ic" if k["sinif"] == IC else "dis"] += k["tutar"]
    return {"kayit": kayit, "ozet": ozet, "turler": turler, "urunler": urunler, "nedenler": nedenler, "matris": matris, "aylar": [{"ay": a, **v} for a, v in sorted(ay.items())], "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    o = r["ozet"]
    _baslik(ws, ["Gösterge", "Değer"], (44, 18))
    for etiket, deger, para_mi in (("Kayıt sayısı", o["kayit"], False), ("Toplam kalitesizlik maliyeti", o["toplam"], True), ("İç hata maliyeti", o["ic"], True), ("Dış hata maliyeti", o["dis"], True), ("Sınıflanmamış", o["diger"], True),
                                   ("İç hata payı %", o["ic_pay"], True), ("Dış hata payı %", o["dis_pay"], True), ("Ciro", o["ciro"], True), ("Kalitesizlik maliyeti ÷ ciro %", o["ciro_oran"], True)):
        _satir(ws, [etiket, deger], para_s=(2,) if para_mi else ())
    ws.append([])
    _baslik(ws, ["Maliyet Türü", "Sınıf", "Kayıt", "Tutar", "Pay %"], dondur="A2")
    for t in r["turler"]:
        _satir(ws, [t["ad"], t["sinif_ad"], t["kayit"], t["tutar"], t["pay"]], para_s=(4, 5), dolgu={2: KIRMIZI if t["sinif_ad"] == DIS else None})
    for h, w in zip("CDE", (10, 18, 10)):
        ws.column_dimensions[h].width = w
    if r["aylar"]:
        ws.append([])
        _baslik(ws, ["Ay", "İç Hata", "Dış Hata", "Toplam"], dondur="A2")
        for a in r["aylar"]:
            _satir(ws, [a["ay"], a["ic"], a["dis"], a["toplam"]], para_s=(2, 3, 4))
    for baslik, alan, ad in (("Ürün Pareto", "urunler", "Ürün"), ("Neden Pareto", "nedenler", "Neden")):
        ws = wb.create_sheet(baslik)
        _baslik(ws, ["Sıra", ad, "Sınıf", "Kayıt", "Tutar", "Pay %", "Kümülatif %", "İç Hata", "Dış Hata"], (6, 34, 7, 8, 16, 9, 12, 16, 16))
        for n, g in enumerate(r[alan], 1):
            _satir(ws, [n, g["ad"], g["sinif"], g["kayit"], g["tutar"], g["pay"], g["kum"], g["ic"] or None, g["dis"] or None], para_s=(5, 6, 7, 8, 9), dolgu={3: KIRMIZI if g["sinif"] == "A" else None})
    ws = wb.create_sheet("Ürün × Tür")
    _baslik(ws, ["Ürün \\ Maliyet Türü"] + [t["ad"] for t in r["turler"]] + ["Toplam"], (30,) + (16,) * (len(r["turler"]) + 1), "B2")
    for m in r["matris"]:
        _satir(ws, [m["urun"]] + [h or None for h in m["hucre"]] + [m["toplam"]], para_s=range(2, len(m["hucre"]) + 3))
    ws = wb.create_sheet("Kayıtlar")
    _baslik(ws, ["Tarih", "Ürün", "Maliyet Türü", "Sınıf", "Neden", "Adet", "Tutar"], (12, 26, 20, 14, 30, 10, 16))
    for k in sorted(r["kayit"], key=lambda k: (k["tarih"] or date.min, katla(k["urun"]))):
        _satir(ws, [k["tarih"], k["urun"], k["tur"], k["sinif"], k["neden"], k["adet"], k["tutar"]], para_s=(7,), tarih_s=(1,))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, ciro: Decimal | None = None, dis_esik: Decimal = Decimal(50)) -> dict:
    kayit, uy = oku(yol)
    r = raporla(kayit, ciro, dis_esik)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Hurda, yeniden işleme, iade ve garanti maliyetlerini ürün ve neden bazında raporlar.")
    p.add_argument("--kayitlar", type=Path, help="Tarih, Ürün, Maliyet Türü, Neden, Adet, Tutar")
    p.add_argument("--ciro", help="Aynı dönemin cirosu; verilirse kalitesizlik maliyetinin ciroya oranı hesaplanır")
    p.add_argument("--dis", default="50", help="Dış hata payı bu yüzdeyi aşarsa uyarılır (varsayılan 50)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kalitesizlik_maliyeti_raporu.xlsx")
    a = p.parse_args(argv)
    ornek = a.kayitlar is None
    yol = ORNEK / "kalitesizlik_kayitlari.csv" if ornek else a.kayitlar
    ciro, dis = para(a.ciro) if a.ciro else Decimal(6500000) if ornek else None, para(a.dis)
    if (a.ciro and (ciro is None or ciro <= 0)) or dis is None or not 0 <= dis <= 100:
        print("[X] --ciro sıfırdan büyük bir tutar, --dis 0 ile 100 arasında yüzde olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, ciro, dis)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {o['kayit']} kayıt · toplam {o['toplam']:,.2f} · iç hata {o['ic']:,.2f} · dış hata {o['dis']:,.2f} (%{o['dis_pay']})" + (f" · ciroya oran %{o['ciro_oran']}" if o["ciro_oran"] is not None else ""))
    print(f"[OK] En maliyetli ürün: {r['urunler'][0]['ad']} (%{r['urunler'][0]['pay']}) · en maliyetli neden: {r['nedenler'][0]['ad']} (%{r['nedenler'][0]['pay']})")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
