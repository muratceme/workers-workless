import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
IC, DIS, DIGER = main.IC, main.DIS, main.DIGER


def k(urun, tur, sinif, neden, tutar, t=date(2026, 7, 1)):
    return {"tarih": t, "urun": urun, "tur": tur, "sinif": sinif, "neden": neden or main.YOK, "adet": None, "tutar": D(tutar)}


def veri():
    return [k("Ü1", "Hurda", IC, "Çapak", 1000), k("ü1", "Yeniden işleme", IC, "çapak", 500, date(2026, 8, 1)), k("Ü2", "İade", DIS, "Etiket", 2000), k("Ü2", "Garanti", DIS, "", 1500, date(2026, 8, 5)), k("Ü3", "Nakliye hasarı", DIGER, "Ambalaj", 250, None)]


class RaporTesti(unittest.TestCase):
    def test_gostergeler(self):
        r = main.raporla(veri(), D(100000))
        self.assertEqual(r["ozet"], {"kayit": 5, "toplam": D(5250), "ic": D(1500), "dis": D(3500), "diger": D(250), "ic_pay": D("28.57"), "dis_pay": D("66.67"), "ciro": D(100000), "ciro_oran": D("5.25")})
        self.assertEqual([(t["ad"], t["sinif_ad"], t["kayit"], t["tutar"], t["pay"]) for t in r["turler"]],
                         [("İade", DIS, 1, D(2000), D("38.10")), ("Garanti", DIS, 1, D(1500), D("28.57")), ("Hurda", IC, 1, D(1000), D("19.05")), ("Yeniden işleme", IC, 1, D(500), D("9.52")), ("Nakliye hasarı", DIGER, 1, D(250), D("4.76"))])
        self.assertEqual([(u["ad"], u["sinif"], u["kayit"], u["tutar"], u["pay"], u["kum"], u["ic"], u["dis"]) for u in r["urunler"]],
                         [("Ü2", "A", 2, D(3500), D("66.67"), D("66.67"), D(0), D(3500)), ("Ü1", "A", 2, D(1500), D("28.57"), D("95.24"), D(1500), D(0)), ("Ü3", "B", 1, D(250), D("4.76"), D("100.00"), D(0), D(0))])
        self.assertEqual([(n["ad"], n["sinif"], n["tutar"], n["kum"]) for n in r["nedenler"]], [("Etiket", "A", D(2000), D("38.10")), ("Çapak", "A", D(1500), D("66.67")), (main.YOK, "A", D(1500), D("95.24")), ("Ambalaj", "B", D(250), D("100.00"))])
        self.assertEqual([(m["urun"], m["hucre"], m["toplam"]) for m in r["matris"]], [("Ü2", [D(2000), D(1500), D(0), D(0), D(0)], D(3500)), ("Ü1", [D(0), D(0), D(1000), D(500), D(0)], D(1500)), ("Ü3", [D(0), D(0), D(0), D(0), D(250)], D(250))])
        self.assertEqual(r["aylar"], [{"ay": "2026-07", "ic": D(1000), "dis": D(2000), "toplam": D(3000)}, {"ay": "2026-08", "ic": D(500), "dis": D(1500), "toplam": D(2000)}])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]), [("Bilgi", "Maliyet türü sınıflanamadı", "Nakliye hasarı"), ("Orta", "Dış hata payı yüksek", "Toplam"), ("Orta", "Nedeni yazılmamış maliyet", main.YOK)])

    def test_esik_ve_sinirlar(self):
        u = lambda **kw: [x["tur"] for x in main.raporla([k("Ü", "Hurda", IC, "N", 100), k("Ü", "İade", DIS, "N", 100)], **kw)["uyarilar"]]   # noqa: E731
        self.assertEqual((u(), u(dis_esik=D("49.99")), main.raporla([k("Ü", "Hurda", IC, "N", 100)])["ozet"]["ciro_oran"]), ([], ["Dış hata payı yüksek"], None))
        for l in ([], [k("Ü", "Hurda", IC, "N", 0)]):
            with self.assertRaises(ValueError):
                main.raporla(l)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "kalitesizlik_kayitlari.csv", cikti, D(6500000))
            o = r["ozet"]
            self.assertEqual((o["kayit"], o["toplam"], o["ic"], o["dis"], o["diger"], o["dis_pay"], o["ciro_oran"]), (11, D(195000), D(92000), D(103000), D(0), D("52.82"), D("3.00")))
            self.assertEqual([(t["ad"], t["tutar"]) for t in r["turler"]], [("İade", D(70500)), ("Hurda", D(69000)), ("Garanti", D(32500)), ("Yeniden işleme", D(15000)), ("Ayıklama", D(8000))])
            self.assertEqual([(u["ad"], u["sinif"], u["tutar"], u["pay"], u["kum"]) for u in r["urunler"]],
                             [("Kapak C", "A", D(75500), D("38.72"), D("38.72")), ("Braket A", "A", D(71500), D("36.67"), D("75.38")), ("Gövde D", "A", D(37500), D("19.23"), D("94.62")), ("Mil E", "B", D(10500), D("5.38"), D("100.00"))])
            self.assertEqual([(n["ad"], n["tutar"]) for n in r["nedenler"]],
                             [("Çapak", D(71500)), ("Boya dökülmesi", D(66500)), ("Ölçü hatası", D(28500)), ("Malzeme hatası", D(15000)), ("Boya kalınlığı", D(9000)), (main.YOK, D(4500))])
            self.assertEqual([(a["ay"], a["toplam"]) for a in r["aylar"]], [("2026-07", D(59000)), ("2026-08", D(70000)), ("2026-09", D(66000))])
            self.assertEqual([(u["onem"], u["tur"]) for u in r["uyarilar"]], [("Orta", "Dış hata payı yüksek"), ("Orta", "Nedeni yazılmamış maliyet")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Ürün Pareto", "Neden Pareto", "Ürün × Tür", "Kayıtlar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "k.csv"
            y.write_text("Kayıt Tarihi;Mamul;Kategori;Hata Nedeni;Miktar;Maliyet\n01.07.2026;Ü;FİRE;Çapak;1.200;1.250,50\n;;Müşteri iadesi;;;500\n;;Rework;;;x\n;;;;;50\n;;Özel;;;10\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["tarih"], x["urun"], x["tur"], x["sinif"], x["neden"], x["adet"], x["tutar"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([(date(2026, 7, 1), "Ü", "Hurda", IC, "Çapak", D(1200), D("1250.50")), (None, "(ürün yok)", "İade", DIS, main.YOK, None, D(500)), (None, "(ürün yok)", "Özel", DIGER, main.YOK, None, D(10))],
                              [("Kayıt okunamadı", "Satır 4"), ("Kayıt okunamadı", "Satır 5")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--kayitlar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--ciro", "1.000.000", "--dis", "40"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--ciro", "0"], ["--ciro", "x"], ["--dis", "101"])], [2] * 3)
            self.assertEqual(main.main(["--kayitlar", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Ürün;Tutar\nA;5\n", "Maliyet Türü;Tutar\n;\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
