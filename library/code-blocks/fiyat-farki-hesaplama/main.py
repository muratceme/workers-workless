"""
Fiyat Farkı Hesaplama — Gerisi Bende kod bloğu
Teknik Ofis › Teknik Ofis Şefi

Sözleşme katsayılarına (a, b1, b2, c...) ve TÜİK endekslerine göre dönem fiyat farkı katsayısını ve tutarını hesaplar.
  Pn = Σ katsayı × (endeksin uygulama ayındaki değeri ÷ endeksin temel aydaki değeri)        (a × In/Io + b1 × Çn/Ço + b2 × Dn/Do + ... + c × Mn/Mo)
  F  = An × B × (Pn − 1)                                                                     (An: uygulama ayında sözleşme fiyatlarıyla yapılan iş tutarı; B: --b, varsayılan 0,90)
  Katsayıların toplamı 1,00 olmalıdır; değilse hesap yapılmaz.
  Pn altı ondalık haneye, fiyat farkı kuruşa yuvarlanır. Pn < 1 ise fiyat farkı eksi çıkar (kesinti).
Katsayılar ve hangi endeksin kullanılacağı sözleşmenizde yazar; temel ayı (--temel-donem) ve endeks değerlerini siz verirsiniz. Kod TÜİK'ten veri indirmez, endeks değeri bilmez.
Bu pakette kaynak teyidi yapılmamıştır. Kapsam dışı: avans, süre uzatımı ve gecikme halindeki endeks seçimi, iş programına göre ödenek dilimleri, güncel fiyatla yapılan işler, ek fiyat farkı düzenlemeleri.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


KATSAYI = {"ad": ("katsayi", "katsayi adi", "sembol"), "deger": ("deger", "agirlik", "oran"), "endeks": ("endeks", "endeks adi", "ilgili endeks")}
ENDEKS = {"donem": ("donem", "ay", "tarih"), "endeks": ("endeks", "endeks adi"), "deger": ("deger", "endeks degeri")}
HAKEDIS = {"donem": ("donem", "uygulama ayi", "ay", "tarih"), "no": ("hakedis no", "hakedis", "no"), "tutar": ("tutar", "an", "hakedis tutari", "yapilan is tutari", "is tutari")}
AYLAR = ("ocak", "subat", "mart", "nisan", "mayis", "haziran", "temmuz", "agustos", "eylul", "ekim", "kasim", "aralik")
Z = Decimal(0)


def donem_adi(x) -> str | None:
    """2026-06 · 06.2026 · 2026/6 · Haziran 2026 · tarih → '2026-06'"""
    if isinstance(x, (date, datetime)):
        return f"{x.year}-{x.month:02d}"
    k = katla(x)
    m = re.fullmatch(r"(\d{4}) (\d{1,2})(?: \d{1,2}.*)?", k)
    if m and 1 <= int(m[2]) <= 12:
        return f"{m[1]}-{int(m[2]):02d}"
    m = re.fullmatch(r"(?:\d{1,2} )?(\d{1,2}) (\d{4})", k)
    if m and 1 <= int(m[1]) <= 12:
        return f"{m[2]}-{int(m[1]):02d}"
    m = re.fullmatch(r"([a-z]+) (\d{4})", k)
    return f"{m[2]}-{AYLAR.index(m[1]) + 1:02d}" if m and m[1] in AYLAR else None


def ondalik(x) -> Decimal | None:
    """Katsayı ve endeks değerleri: 0,25 · 0.25 · 1.234,56 (nokta yalnız virgülle birlikte binliktir)"""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    s = s.replace(".", "").replace(",", ".") if "," in s and "." in s else s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def oku(katsayi_yol: Path, endeks_yol: Path, hakedis_yol: Path) -> tuple[list[dict], dict, list[dict], list[dict]]:
    uy, katsayilar, endeksler, hakedisler = [], [], {}, []
    for r in kayitlar(katsayi_yol, KATSAYI, ("ad", "deger", "endeks")):
        ad, d, e = metin(r["ad"]), ondalik(r["deger"]), metin(r["endeks"])
        if not ad:
            continue
        if d is None or d < 0 or not e:
            raise ValueError(f"{katsayi_yol.name} satır {r['_satir']}: '{ad}' katsayısının değeri veya endeksi okunamadı.")
        if d:
            katsayilar.append({"ad": ad, "deger": d, "endeks": e})
    for r in kayitlar(endeks_yol, ENDEKS, ("donem", "endeks", "deger")):
        e = metin(r["endeks"])
        if not e:
            continue
        dn, v = donem_adi(r["donem"]), ondalik(r["deger"])
        if dn is None or v is None or v <= 0:
            uy.append(uyari("Yüksek", "Endeks satırı okunamadı", f"{e} ({endeks_yol.name} satır {r['_satir']})", f"Dönem '{metin(r['donem'])}', değer '{metin(r['deger'])}'. Satır alınmadı."))
            continue
        if (dn, katla(e)) in endeksler and endeksler[dn, katla(e)] != v:
            uy.append(uyari("Yüksek", "Endeks iki farklı değerle girilmiş", f"{e} · {dn}", f"{endeksler[dn, katla(e)]:f} ve {v:f}. İlki kullanıldı."))
            continue
        endeksler[dn, katla(e)] = v
    for r in kayitlar(hakedis_yol, HAKEDIS, ("donem", "tutar")):
        if not metin(r["donem"]) and not metin(r["tutar"]):
            continue
        dn, t = donem_adi(r["donem"]), para(r["tutar"])
        if dn is None or t is None:
            uy.append(uyari("Yüksek", "Hakediş satırı okunamadı", f"{hakedis_yol.name} satır {r['_satir']}", f"Dönem '{metin(r['donem'])}', tutar '{metin(r['tutar'])}'. Satır alınmadı."))
            continue
        hakedisler.append({"no": metin(r["no"]), "donem": dn, "an": t})
    return katsayilar, endeksler, hakedisler, uy


def hesapla(katsayilar: list[dict], endeksler: dict, hakedisler: list[dict], temel: str, b: Decimal = Decimal("0.90")) -> dict:
    if not katsayilar:
        raise ValueError("Katsayı dosyasında değeri sıfırdan büyük katsayı yok.")
    if not hakedisler:
        raise ValueError("Hesaplanacak hakediş yok.")
    toplam = sum((k["deger"] for k in katsayilar), Z)
    if abs(toplam - 1) > Decimal("0.00001"):
        raise ValueError(f"Katsayıların toplamı 1,00 olmalı; dosyadaki toplam {toplam:f}. Sözleşmedeki katsayıları kontrol edin.")
    eksik = [k["endeks"] for k in katsayilar if (temel, katla(k["endeks"])) not in endeksler]
    if eksik:
        raise ValueError(f"Temel ay ({temel}) için endeks değeri yok: {', '.join(eksik)}.")
    for h in hakedisler:
        h["detay"], h["eksik"] = [], [k["endeks"] for k in katsayilar if (h["donem"], katla(k["endeks"])) not in endeksler]
        if h["eksik"]:
            h["pn"] = h["fark"] = None
            continue
        pn = Z
        for k in katsayilar:
            n, o = endeksler[h["donem"], katla(k["endeks"])], endeksler[temel, katla(k["endeks"])]
            pn += k["deger"] * n / o
            h["detay"].append({"ad": k["ad"], "endeks": k["endeks"], "deger": k["deger"], "o": o, "n": n, "oran": yuvarla(n / o, "0.000001"), "katki": yuvarla(k["deger"] * n / o, "0.000001")})
        h["pn"] = yuvarla(pn, "0.000001")
        h["fark"] = yuvarla(h["an"] * b * (h["pn"] - 1))
    hesaplanan = [h for h in hakedisler if h["fark"] is not None]
    kum = Z
    for h in sorted(hesaplanan, key=lambda h: h["donem"]):
        kum += h["fark"]
        h["kumulatif"] = kum
    return {"hakedisler": sorted(hakedisler, key=lambda h: (h["donem"], h["no"])), "katsayilar": katsayilar, "temel": temel, "b": b,
            "toplam": {"an": sum((h["an"] for h in hesaplanan), Z), "fark": kum, "hesaplanan": len(hesaplanan), "hesaplanamayan": len(hakedisler) - len(hesaplanan)}}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    t = r["toplam"]
    ws = wb.active
    ws.title = "Fiyat Farkı"
    _baslik(ws, ["Hakediş No", "Uygulama Ayı", "An (yapılan iş tutarı)", "Pn", "Pn − 1", "B", "Fiyat Farkı (F)", "Kümülatif Fiyat Farkı", "Not"], (12, 12, 18, 12, 12, 7, 16, 18, 50))
    for h in r["hakedisler"]:
        n = _satir(ws, [h["no"], h["donem"], h["an"], h["pn"], h["pn"] - 1 if h["pn"] is not None else None, r["b"], h["fark"], h.get("kumulatif"),
                        f"Hesaplanamadı: {h['donem']} için endeks yok ({', '.join(h['eksik'])})" if h["eksik"] else ""], para_s=(3, 7, 8),
                   dolgu={7: None if h["fark"] is None else (KIRMIZI if h["fark"] < 0 else YESIL), 9: KIRMIZI if h["eksik"] else None}, sar=(9,))
        ws.cell(n, 4).number_format = ws.cell(n, 5).number_format = "0.000000"
    _satir(ws, ["Toplam", "", t["an"], None, None, None, t["fark"]], para_s=(3, 7))
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.append([])
    ws.append([f"F = An × B × (Pn − 1); temel ay {r['temel']}. Katsayılar ve endeksler dosyalarınızdan alınmıştır. Avans, süre uzatımı, gecikme ve ek fiyat farkı düzenlemeleri hesaba katılmaz; "
               "idarenin hesabı esastır."])
    ws = wb.create_sheet("Pn Detayı")
    _baslik(ws, ["Uygulama Ayı", "Katsayı", "Katsayı Değeri", "Endeks", f"Temel Ay Değeri ({r['temel']})", "Uygulama Ayı Değeri", "Endeks Oranı", "Pn'e Katkı"], (12, 9, 10, 34, 16, 16, 12, 12))
    for h in r["hakedisler"]:
        for d in h["detay"]:
            n = _satir(ws, [h["donem"], d["ad"], d["deger"], d["endeks"], d["o"], d["n"], d["oran"], d["katki"]], para_s=(5, 6))
            ws.cell(n, 7).number_format = ws.cell(n, 8).number_format = "0.000000"
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(katsayi_yol: Path, endeks_yol: Path, hakedis_yol: Path, cikti: Path, temel: str, b: Decimal = Decimal("0.90")) -> dict:
    katsayilar, endeksler, hakedisler, uy = oku(katsayi_yol, endeks_yol, hakedis_yol)
    r = hesapla(katsayilar, endeksler, hakedisler, temel, b)
    for h in r["hakedisler"]:
        if h["eksik"]:
            uy.append(uyari("Yüksek", "Fiyat farkı hesaplanamadı", f"{h['no'] or 'Hakediş'} · {h['donem']}", f"Bu ay için endeks değeri yok: {', '.join(h['eksik'])}. Endeks açıklanınca dosyaya ekleyip yeniden çalıştırın."))
        elif h["donem"] <= temel:
            uy.append(uyari("Orta", "Uygulama ayı temel aydan önce veya aynı", f"{h['no'] or 'Hakediş'} · {h['donem']}", f"Temel ay {temel}. Dönemleri kontrol edin."))
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Sözleşme katsayıları ve endekslerden fiyat farkı katsayısını (Pn) ve tutarını hesaplar.")
    p.add_argument("--katsayilar", type=Path, default=ORNEK / "katsayilar.csv", help="Katsayı (a, b1, ...), Değer, Endeks (endeks dosyasındaki ad)")
    p.add_argument("--endeksler", type=Path, default=ORNEK / "endeksler.csv", help="Dönem, Endeks, Değer")
    p.add_argument("--hakedisler", type=Path, default=ORNEK / "hakedisler.csv", help="Hakediş No, Dönem (uygulama ayı), Tutar (An)")
    p.add_argument("--temel-donem", help="Temel endekslerin ait olduğu ay, YYYY-AA (sözleşmenize göre; örnek veride 2026-03)")
    p.add_argument("--b", default="0,90", help="B katsayısı (varsayılan 0,90)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "fiyat_farki.xlsx")
    a = p.parse_args(argv)
    temel = donem_adi(a.temel_donem or ("2026-03" if a.katsayilar == ORNEK / "katsayilar.csv" else ""))
    b = ondalik(a.b)
    if temel is None or b is None or not 0 < b <= 1:
        print("[X] --temel-donem YYYY-AA biçiminde verilmeli (ör. 2026-03); --b 0 ile 1 arasında olmalı")
        return 2
    for y in (a.katsayilar, a.endeksler, a.hakedisler):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.katsayilar, a.endeksler, a.hakedisler, a.cikti, temel, b)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    for h in r["hakedisler"]:
        print(f"[!] {h['no'] or 'Hakediş'} · {h['donem']}: hesaplanamadı (endeks yok)" if h["eksik"] else f"[OK] {h['no'] or 'Hakediş'} · {h['donem']}: An {h['an']:,.2f} · Pn {h['pn']} · fiyat farkı {h['fark']:,.2f}")
    t = r["toplam"]
    print(f"[OK] {t['hesaplanan']} hakediş · toplam fiyat farkı {t['fark']:,.2f} · rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
