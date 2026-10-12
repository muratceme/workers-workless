import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
KATSAYI = [{"ad": "a", "deger": D("0.30"), "endeks": "İşçilik"}, {"ad": "b1", "deger": D("0.25"), "endeks": "Çimento"}, {"ad": "b2", "deger": D("0.25"), "endeks": "Demir"}, {"ad": "c", "deger": D("0.20"), "endeks": "Makine"}]


def en(**donemler):
    """en(d03=(1000, 2000, 3000, 1500)) → 2026-03 için işçilik, çimento, demir, makine"""
    return {(f"2026-{d[1:]}", ad): D(v) for d, degerler in donemler.items() for ad, v in zip(("iscilik", "cimento", "demir", "makine"), degerler)}


ENDEKS = en(d03=(1000, 2000, 3000, 1500), d06=(1100, 2100, 3300, 1500), d07=(1100, 2200, 3600, 1575), d04=(900, 2000, 2700, 1500))


def hk(donem, an, no=""):
    return {"no": no, "donem": f"2026-{donem:02d}", "an": D(an)}


class HesapTesti(unittest.TestCase):
    def test_pn_ve_fiyat_farki(self):
        r = main.hesapla(KATSAYI, ENDEKS, [hk(6, 2000000, "1"), hk(7, 3000000, "2")], "2026-03")
        self.assertEqual([(h["pn"], h["fark"], h["kumulatif"]) for h in r["hakedisler"]], [(D("1.067500"), D("121500.00"), D("121500.00")), (D("1.115000"), D("310500.00"), D("432000.00"))])
        self.assertEqual([(d["ad"], d["oran"], d["katki"]) for d in r["hakedisler"][0]["detay"]], [("a", D("1.100000"), D("0.330000")), ("b1", D("1.050000"), D("0.262500")), ("b2", D("1.100000"), D("0.275000")), ("c", D("1.000000"), D("0.200000"))])
        self.assertEqual(r["toplam"], {"an": D(5000000), "fark": D("432000.00"), "hesaplanan": 2, "hesaplanamayan": 0})
        h = main.hesapla(KATSAYI, ENDEKS, [hk(4, 1000000)], "2026-03")["hakedisler"][0]                    # endeksler düşmüş: 0,27 + 0,25 + 0,225 + 0,20 = 0,945
        self.assertEqual((h["pn"], h["fark"]), (D("0.945000"), D("-49500.00")))
        self.assertEqual(main.hesapla(KATSAYI, ENDEKS, [hk(6, 2000000)], "2026-03", D(1))["hakedisler"][0]["fark"], D("135000.00"))
        self.assertEqual(main.hesapla(KATSAYI, ENDEKS, [hk(3, 500)], "2026-03")["hakedisler"][0]["fark"], D("0.00"))
        h = main.hesapla([{"ad": "a", "deger": D(1), "endeks": "İşçilik"}], en(d03=(3,), d06=(4,)), [hk(6, 1000)], "2026-03")["hakedisler"][0]
        self.assertEqual((h["pn"], h["fark"]), (D("1.333333"), D("300.00")))                                 # 1000 × 0,9 × 0,333333 = 299,9997

    def test_eksik_veri_ve_hatalar(self):
        r = main.hesapla(KATSAYI, ENDEKS, [hk(6, 100), hk(8, 100)], "2026-03")
        self.assertEqual(([(h["donem"], h["pn"], h["eksik"]) for h in r["hakedisler"]][1], r["toplam"]["hesaplanamayan"]), (("2026-08", None, ["İşçilik", "Çimento", "Demir", "Makine"]), 1))
        for k, e, h, t in ((KATSAYI[:3], ENDEKS, [hk(6, 1)], "2026-03"), (KATSAYI, ENDEKS, [hk(6, 1)], "2026-01"), ([], ENDEKS, [hk(6, 1)], "2026-03"), (KATSAYI, ENDEKS, [], "2026-03")):
            with self.assertRaises(ValueError):
                main.hesapla(k, e, h, t)
        self.assertEqual([main.donem_adi(x) for x in ("2026-06", "06.2026", "2026/6", "Haziran 2026", "15.06.2026", date(2026, 6, 1), "2026-13", "yaz")], ["2026-06"] * 6 + [None, None])
        self.assertEqual([main.ondalik(x) for x in ("0,25", "0.25", "1.234,56", "1100.00", "x")], [D("0.25"), D("0.25"), D("1234.56"), D("1100.00"), None])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "f.xlsx"
            r = main.calistir(ORNEK / "katsayilar.csv", ORNEK / "endeksler.csv", ORNEK / "hakedisler.csv", cikti, "2026-03")
            self.assertEqual(([(h["no"], h["pn"], h["fark"]) for h in r["hakedisler"]], r["toplam"]["fark"], [u["tur"] for u in r["uyarilar"]]),
                             ([("1", D("1.067500"), D("121500.00")), ("2", D("1.115000"), D("310500.00")), ("3", None, None)], D("432000.00"), ["Fiyat farkı hesaplanamadı"]))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Fiyat Farkı", "Pn Detayı", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            k, e, h = Path(tmp) / "k.csv", Path(tmp) / "e.csv", Path(tmp) / "h.csv"
            k.write_text("Katsayı;Değer;Endeks\na;0,60;İşçilik\nb1;0.40;Demir\nb2;0;Çimento\n", encoding="utf-8")
            e.write_text("Dönem;Endeks;Değer\nMart 2026;İşçilik;1.000,00\n03.2026;DEMİR;200\n2026-06;İşçilik;1100\n2026-06;Demir;220\n2026-06;Demir;230\nyaz;Demir;1\n2026-06;Demir;sıfır\n", encoding="utf-8")
            h.write_text("Dönem;Tutar\nHaziran 2026;1.000.000,00\n2026-02;500\nbilinmiyor;5\n", encoding="utf-8")
            kl, el, hl, uy = main.oku(k, e, h)
            self.assertEqual(([x["ad"] for x in kl], el, [(x["donem"], x["an"]) for x in hl]),
                             (["a", "b1"], {("2026-03", "iscilik"): D("1000.00"), ("2026-03", "demir"): D(200), ("2026-06", "iscilik"): D(1100), ("2026-06", "demir"): D(220)},
                              [("2026-06", D("1000000.00")), ("2026-02", D(500))]))
            self.assertEqual(sorted(u["tur"] for u in uy), ["Endeks iki farklı değerle girilmiş", "Endeks satırı okunamadı", "Endeks satırı okunamadı", "Hakediş satırı okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            r = main.calistir(k, e, h, Path(c), "2026-03")
            self.assertEqual(([(x["donem"], x["pn"], x["fark"]) for x in r["hakedisler"]], "Fiyat farkı hesaplanamadı" in [u["tur"] for u in r["uyarilar"]]),
                             ([("2026-02", None, None), ("2026-06", D("1.100000"), D("90000.00"))], True))
            arg = ["--katsayilar", str(k), "--endeksler", str(e), "--hakedisler", str(h), "--cikti", c]
            self.assertEqual(main.main(arg + ["--temel-donem", "Mart 2026", "--b", "1"]), 0)
            self.assertEqual(main.main(arg), 2)                                              # kendi verinizle temel ay zorunlu
            self.assertEqual(main.main(arg + ["--temel-donem", "2026-03", "--b", "1,5"]), 2)
            self.assertEqual(main.main(arg + ["--temel-donem", "2025-01"]), 1)
            k.write_text("Katsayı;Değer;Endeks\na;0,60;İşçilik\nb1;0,30;Demir\n", encoding="utf-8")
            self.assertEqual(main.main(arg + ["--temel-donem", "2026-03"]), 1)              # katsayılar toplamı 0,90
            self.assertEqual(main.main(arg[:1] + [str(Path(tmp) / "yok.csv")] + arg[2:] + ["--temel-donem", "2026-03"]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
