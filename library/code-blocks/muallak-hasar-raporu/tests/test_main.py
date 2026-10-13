import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)


def a(no, yas, ilk, guncel, odenen=0, brans="Kasko"):
    return {"no": no, "brans": brans, "ihbar": BUGUN - timedelta(days=yas), "kapanis": None, "ilk": D(ilk), "guncel": D(guncel), "odenen": D(odenen)}


def k(no, ilk, odenen, sure=30, brans="Kasko", guncel=0):
    return {"no": no, "brans": brans, "ihbar": BUGUN - timedelta(days=sure + 10), "kapanis": BUGUN - timedelta(days=10), "ilk": D(ilk), "guncel": D(guncel), "odenen": D(odenen)}


class RaporTesti(unittest.TestCase):
    def test_acik_ve_kapali(self):
        l = [a("a1", 50, 100, 120), a("a2", 200, 200, 150, 100), a("a3", 400, 1000, 1000, brans="Yangın"), k("k1", 100, 130), k("k2", 100, 120, 40), k("k3", 100, 0, 50), k("k4", 0, 50, 60, "Yangın"), k("k5", 100, 80, 80)]
        r = main.raporla(l, BUGUN, dilimler=(D(150), D(1000)))
        self.assertEqual([(b["brans"], b["acik"]["dosya"], b["acik"]["muallak"], b["acik"]["pay"], b["acik"]["ort"], b["acik"]["enb"], b["acik"]["ilk"], b["acik"]["nihai"], b["acik"]["gelisim"], b["acik"]["ort_yas"], b["yas"]) for b in r["branslar"]],
                         [("Yangın", 1, D(1000), D("78.74"), D("1000.00"), D(1000), D(1000), D(1000), D("0.00"), D("400.00"), [D(0), D(0), D(0), D(1000)]),
                          ("Kasko", 2, D(270), D("21.26"), D("135.00"), D(150), D(300), D(370), D("23.33"), D("125.00"), [D(120), D(0), D(150), D(0)])])
        self.assertEqual([(b["brans"], b["kapali"]) for b in r["branslar"]],
                         [("Yangın", {"dosya": 1, "ilk": D(0), "odenen": D(0), "oran": None, "yetersiz": 0, "fazla": 0, "odemesiz": 0, "ilksiz": 1, "ort_sure": D("60.00")}),
                          ("Kasko", {"dosya": 4, "ilk": D(400), "odenen": D(330), "oran": D("82.50"), "yetersiz": 1, "fazla": 1, "odemesiz": 1, "ilksiz": 0, "ort_sure": D("50.00")})])
        self.assertEqual({d["no"]: d["sonuc"] for d in r["kapali"]}, {"k1": "Yetersiz", "k2": "Uyumlu", "k3": "Fazla", "k4": "İlk rezerv yok", "k5": "Uyumlu"})
        self.assertEqual([(y["dilim"], y["dosya"], y["tutar"], y["pay"]) for y in r["yaslar"]], [("0–90", 1, D(120), D("9.45")), ("91–180", 0, D(0), D("0.00")), ("181–365", 1, D(150), D("11.81")), ("365+", 1, D(1000), D("78.74"))])
        self.assertEqual([(y["dilim"], y["dosya"], y["tutar"]) for y in r["dilimler"]], [("150 altı", 1, D(120)), ("150 – 1.000", 1, D(150)), ("1.000 ve üzeri", 1, D(1000))])
        self.assertEqual((r["acik_ozet"]["dosya"], r["acik_ozet"]["muallak"], r["acik_ozet"]["gelisim"], r["kapali_ozet"]["dosya"], r["kapali_ozet"]["oran"], r["kapali_ozet"]["ilksiz"], [x["no"] for x in r["acik"]], r["uyarilar"]),
                         (3, D(1270), D("5.38"), 5, D("82.50"), 1, ["a3", "a2", "a1"], []))

    def test_uyarilar(self):
        l = [a("a1", 50, 100, 120), a("a2", 200, 200, 150, 100), k("k1", 100, 130), k("k2", 100, 120), k("k3", 100, 0), k("k5", 100, 80)]
        u = lambda **kw: sorted((x["onem"], x["tur"]) for x in main.raporla([dict(d) for d in l], BUGUN, **kw)["uyarilar"])   # noqa: E731
        self.assertEqual((u(), u(asgari=2), u(sapma=D(10)), u(sapma=D(10), asgari=5)), ([], [("Orta", "Açık dosyalarda rezerv artışı")], [("Orta", "İlk rezerv fazla ayrılmış")], []))
        r = main.raporla([k("k1", 100, 150), k("k2", 100, 121, guncel=5), k("k3", 100, 100), a("x", 91, 0, 50), a("y", 90, 0, 0), a("z", 365, 0, 0), a("w", 181, 0, 0)], BUGUN)
        self.assertEqual(([(x["onem"], x["tur"], x["kim"]) for x in r["uyarilar"]], r["kapali_ozet"]["oran"], r["acik_ozet"]["gelisim"], {d["no"]: d["yas_dilim"] for d in r["acik"]}),
                         ([("Orta", "Kapalı dosyada muallak kalmış", "k2"), ("Yüksek", "İlk rezerv yetersiz kalmış", "Kasko")], D("123.67"), None, {"x": "91–180", "y": "0–90", "z": "181–365", "w": "181–365"}))
        self.assertEqual(main.raporla([k("k1", 100, 100)], BUGUN)["acik_ozet"], {"dosya": 0, "muallak": D(0), "ort": None, "enb": None, "pay": None, "ilk": D(0), "nihai": D(0), "gelisim": None, "ort_yas": None})
        with self.assertRaises(ValueError):
            main.raporla([], BUGUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "hasar_dosyalari.csv", cikti, BUGUN)
            self.assertEqual([(b["brans"], b["acik"]["dosya"], b["acik"]["muallak"], b["acik"]["gelisim"], b["kapali"]["dosya"], b["kapali"]["ilk"], b["kapali"]["odenen"], b["kapali"]["oran"], b["kapali"]["yetersiz"], b["kapali"]["fazla"])
                              for b in r["branslar"]],
                             [("Trafik", 2, D(670000), D("48.89"), 3, D(380000), D(531000), D("139.74"), 2, 0), ("Yangın", 1, D(650000), D("6.25"), 1, D(400000), D(380000), D("95.00"), 0, 0),
                              ("Kasko", 3, D(210000), D("34.21"), 3, D(150000), D(88000), D("58.67"), 1, 1)])
            self.assertEqual((r["acik_ozet"]["dosya"], r["acik_ozet"]["muallak"], r["kapali_ozet"]["dosya"], r["kapali_ozet"]["ilk"], r["kapali_ozet"]["odenen"], r["kapali_ozet"]["oran"]), (6, D(1530000), 7, D(930000), D(999000), D("107.42")))
            self.assertEqual({d["no"]: d["yas"] for d in r["acik"]}, {"M-01": 22, "M-02": 133, "M-03": 275, "M-04": 576, "M-05": 68, "M-06": 183})
            self.assertEqual([(y["dilim"], y["dosya"], y["tutar"]) for y in r["yaslar"]], [("0–90", 2, D(190000)), ("91–180", 1, D(140000)), ("181–365", 2, D(680000)), ("365+", 1, D(520000))])
            self.assertEqual([(y["dilim"], y["dosya"]) for y in r["dilimler"]], [("50.000 altı", 2), ("50.000 – 250.000", 2), ("250.000 – 1.000.000", 2), ("1.000.000 ve üzeri", 0)])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]], [("Yüksek", "İlk rezerv yetersiz kalmış", "Trafik"), ("Orta", "Açık dosyalarda rezerv artışı", "Kasko"), ("Orta", "İlk rezerv fazla ayrılmış", "Kasko")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Branşlar (Açık)", "Rezerv Yeterliliği", "Açık Dosyalar", "Kapanan Dosyalar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "d.csv"
            y.write_text("Hasar No;Ürün;Açılış Tarihi;Kapanma Tarihi;İlk Rezerv;Muallak;Ödenen\nD1;Kasko;01.09.2026;;1.500,50;2.000;\nD2;;02.09.2026;05.10.2026;100;;90\nd1;;03.09.2026;;;;\nD3;;x;;;;\nD4;;03.09.2026;x;;;\nD5;;03.09.2026;;-1;;\n",
                         encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["brans"], x["ihbar"], x["kapanis"], x["ilk"], x["guncel"], x["odenen"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("D1", "Kasko", date(2026, 9, 1), None, D("1500.50"), D(2000), D(0)), ("D2", "(branş yok)", date(2026, 9, 2), date(2026, 10, 5), D(100), D(0), D(90))],
                              [("Dosya okunamadı", "d1"), ("Dosya okunamadı", "D3"), ("Dosya okunamadı", "D4"), ("Dosya okunamadı", "D5")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--dosyalar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "12.10.2026", "--sapma", "15", "--min", "1", "--dilimler", "1000;5000"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--sapma", "0"], ["--sapma", "100"], ["--min", "0"], ["--dilimler", "5;1"])], [2] * 5)
            self.assertEqual(main.main(["--dosyalar", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Dosya No;Branş\nA;B\n", "Dosya No;İhbar Tarihi\n;01.03.2026\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
