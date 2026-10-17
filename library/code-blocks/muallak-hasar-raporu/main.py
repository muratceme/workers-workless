"""
Muallak Hasar Raporu — Gerisi Bende kod bloğu
Hasar Yönetimi › Hasar Müdürü

Açık dosyaların branş, yaş ve tutar dağılımını; kapanan dosyalardan rezerv yeterliliği göstergelerini raporlar.
  Açık dosya            = kapanış tarihi boş;   yaş (gün) = durum tarihi − ihbar tarihi;   dilimler 0–90, 91–180, 181–365, 365+
  Tahmini nihai hasar   = güncel muallak + bugüne kadar ödenen             (açık dosya)
  Rezerv gelişimi %     = (tahmini nihai hasar − ilk muallak) ÷ ilk muallak × 100   (açık dosyalar, toplam üzerinden)
  Gerçekleşme oranı %   = ödenen ÷ ilk muallak × 100                        (kapalı dosyalar, toplam üzerinden)
                          %100'ün üstü: ilk rezerv yetersiz kalmış;   altı: fazla ayrılmış
  Yetersiz / fazla rezervli dosya = ödeneni ilk muallağın (1 ± --sapma %20) dışında kalan kapalı dosya
  Uyarı                 = en az --min (3) dosyası olan branşta oran sapma bandının dışındaysa
Aktüeryal karşılık hesabı değildir; IBNR ve hasar gelişim üçgenleri kapsam dışıdır.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"no": ("dosya no", "hasar dosya no", "hasar no"), "brans": ("brans", "urun"), "ihbar": ("ihbar tarihi", "dosya acilis tarihi", "acilis tarihi"), "kapanis": ("kapanis tarihi", "kapanma tarihi"),
         "ilk": ("ilk muallak", "ilk rezerv", "acilis rezervi"), "guncel": ("guncel muallak", "muallak tutar", "muallak", "guncel rezerv"), "odenen": ("odenen tutar", "odenen", "odenen hasar")}
YAS = ((90, "0–90"), (180, "91–180"), (365, "181–365"), (None, "365+"))


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("no", "ihbar")):
        no = metin(r["no"])
        if not no:
            continue
        ihbar, t = tarih(r["ihbar"]), {a: para(r[a]) if metin(r[a]) else Decimal(0) for a in ("ilk", "guncel", "odenen")}
        kapanis = tarih(r["kapanis"])
        if ihbar is None or any(v is None or v < 0 for v in t.values()) or katla(no) in gorulen or (metin(r["kapanis"]) and kapanis is None):
            uy.append(uyari("Yüksek", "Dosya okunamadı", no, f"Satır {r['_satir']}: " + ("aynı dosya numarası daha önce geçti" if katla(no) in gorulen else "ihbar tarihi, kapanış tarihi ya da tutarlardan biri okunamadı") + ". Satır alınmadı."))
            continue
        gorulen.add(katla(no))
        l.append({"no": no, "brans": metin(r["brans"]) or "(branş yok)", "ihbar": ihbar, "kapanis": kapanis, **t})
    return l, uy


def _oran(pay: Decimal, payda: Decimal) -> Decimal | None:
    return yuvarla(Decimal(pay) * 100 / Decimal(payda)) if payda else None


def _b(x: Decimal) -> str:
    return f"{int(x):,}".replace(",", ".")


def dilim_adlari(dilimler: tuple[Decimal, ...]) -> list[str]:
    return [f"{_b(dilimler[0])} altı"] + [f"{_b(a)} – {_b(u)}" for a, u in zip(dilimler, dilimler[1:])] + [f"{_b(dilimler[-1])} ve üzeri"]


def raporla(dosyalar: list[dict], bugun: date, sapma: Decimal = Decimal(20), asgari: int = 3, dilimler: tuple[Decimal, ...] = (Decimal(50000), Decimal(250000), Decimal(1000000))) -> dict:
    if not dosyalar:
        raise ValueError("Raporlanacak hasar dosyası yok.")
    uy, adlar = [], dilim_adlari(dilimler)
    ust, alt = (100 + sapma) / 100, (100 - sapma) / 100
    for d in dosyalar:
        d["acik"] = d["kapanis"] is None
        if d["acik"]:
            d["yas"] = max((bugun - d["ihbar"]).days, 0)
            d["yas_dilim"] = next(e for s, e in YAS if s is None or d["yas"] <= s)
            d["dilim"] = adlar[sum(1 for x in dilimler if d["guncel"] >= x)]
            d["nihai"] = d["guncel"] + d["odenen"]
            d["sonuc"] = None
        else:
            d["sure"] = max((d["kapanis"] - d["ihbar"]).days, 0)
            d["sonuc"] = "İlk rezerv yok" if d["ilk"] == 0 else "Yetersiz" if d["odenen"] > d["ilk"] * ust else "Fazla" if d["odenen"] < d["ilk"] * alt else "Uyumlu"
            if d["guncel"] > 0:
                uy.append(uyari("Orta", "Kapalı dosyada muallak kalmış", d["no"], f"Dosya {d['kapanis']:%d.%m.%Y} tarihinde kapanmış ama güncel muallak {d['guncel']:,.2f}. Rezervi serbest bırakın ya da kapanışı kontrol edin."))
    acik, kapali = [d for d in dosyalar if d["acik"]], [d for d in dosyalar if not d["acik"]]
    top_muallak = sum((d["guncel"] for d in acik), Decimal(0))

    def acik_ozet(l: list[dict]) -> dict:
        ilkli = [d for d in l if d["ilk"] > 0]
        ilk, nihai, tutar = sum((d["ilk"] for d in ilkli), Decimal(0)), sum((d["nihai"] for d in ilkli), Decimal(0)), sum((d["guncel"] for d in l), Decimal(0))
        return {"dosya": len(l), "muallak": tutar, "ort": yuvarla(tutar / len(l)) if l else None, "enb": max((d["guncel"] for d in l), default=None), "pay": _oran(tutar, top_muallak), "ilk": ilk, "nihai": nihai,
                "gelisim": _oran(nihai - ilk, ilk), "ort_yas": yuvarla(Decimal(sum(d["yas"] for d in l)) / len(l)) if l else None}

    def kapali_ozet(l: list[dict]) -> dict:
        ilkli = [d for d in l if d["ilk"] > 0]
        ilk, odenen = sum((d["ilk"] for d in ilkli), Decimal(0)), sum((d["odenen"] for d in ilkli), Decimal(0))
        return {"dosya": len(l), "ilk": ilk, "odenen": odenen, "oran": _oran(odenen, ilk), "yetersiz": sum(1 for d in l if d["sonuc"] == "Yetersiz"), "fazla": sum(1 for d in l if d["sonuc"] == "Fazla"),
                "odemesiz": sum(1 for d in l if d["odenen"] == 0), "ilksiz": len(l) - len(ilkli), "ort_sure": yuvarla(Decimal(sum(d["sure"] for d in l)) / len(l)) if l else None}

    grup = {}
    for d in dosyalar:
        grup.setdefault(katla(d["brans"]), []).append(d)
    branslar = []
    for g in grup.values():
        a, k = [d for d in g if d["acik"]], [d for d in g if not d["acik"]]
        b = {"brans": g[0]["brans"], "acik": acik_ozet(a), "kapali": kapali_ozet(k), "yas": [sum((d["guncel"] for d in a if d["yas_dilim"] == e), Decimal(0)) for _, e in YAS]}
        branslar.append(b)
        ko, ao = b["kapali"], b["acik"]
        if ko["dosya"] >= asgari and ko["oran"] is not None and ko["oran"] > 100 + sapma:
            uy.append(uyari("Yüksek", "İlk rezerv yetersiz kalmış", b["brans"], f"Kapanan {ko['dosya']} dosyada ödenen {ko['odenen']:,.2f}, ilk muallak {ko['ilk']:,.2f} (gerçekleşme %{ko['oran']}). {ko['yetersiz']} dosyada ödeme ilk rezervi "
                            f"%{sapma}'den fazla aşmış. Açık dosyaların ilk rezerv yaklaşımını gözden geçirin."))
        elif ko["dosya"] >= asgari and ko["oran"] is not None and ko["oran"] < 100 - sapma:
            uy.append(uyari("Orta", "İlk rezerv fazla ayrılmış", b["brans"], f"Kapanan {ko['dosya']} dosyada ödenen {ko['odenen']:,.2f}, ilk muallak {ko['ilk']:,.2f} (gerçekleşme %{ko['oran']}). {ko['odemesiz']} dosya ödemesiz kapanmış."))
        if ao["dosya"] >= asgari and ao["gelisim"] is not None and ao["gelisim"] > sapma:
            uy.append(uyari("Orta", "Açık dosyalarda rezerv artışı", b["brans"], f"{ao['dosya']} açık dosyada ilk muallak {ao['ilk']:,.2f}, tahmini nihai hasar {ao['nihai']:,.2f} (gelişim %{ao['gelisim']})."))
    branslar.sort(key=lambda b: (-b["acik"]["muallak"], katla(b["brans"])))
    return {"dosyalar": dosyalar, "acik": sorted(acik, key=lambda d: (-d["guncel"], katla(d["no"]))), "kapali": sorted(kapali, key=lambda d: (d["kapanis"], katla(d["no"]))), "branslar": branslar,
            "acik_ozet": acik_ozet(acik), "kapali_ozet": kapali_ozet(kapali),
            "yaslar": [{"dilim": e, "dosya": sum(1 for d in acik if d["yas_dilim"] == e), "tutar": (t := sum((d["guncel"] for d in acik if d["yas_dilim"] == e), Decimal(0))), "pay": _oran(t, top_muallak)} for _, e in YAS],
            "dilimler": [{"dilim": e, "dosya": sum(1 for d in acik if d["dilim"] == e), "tutar": (t := sum((d["guncel"] for d in acik if d["dilim"] == e), Decimal(0))), "pay": _oran(t, top_muallak)} for e in adlar], "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, bugun: date, sapma: Decimal) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    a, k = r["acik_ozet"], r["kapali_ozet"]
    _baslik(ws, ["Gösterge", "Değer"], (48, 18))
    for etiket, deger, para_mi in (("Açık dosya", a["dosya"], False), ("Toplam muallak", a["muallak"], True), ("Dosya başına ortalama muallak", a["ort"], True), ("En büyük dosya", a["enb"], True),
                                   ("Açık dosyaların ortalama yaşı (gün)", a["ort_yas"], True), ("Açık dosyalarda rezerv gelişimi %", a["gelisim"], True), ("Kapanan dosya", k["dosya"], False),
                                   ("Kapanan dosyalarda ilk muallak", k["ilk"], True), ("Kapanan dosyalarda ödenen", k["odenen"], True), ("Gerçekleşme oranı % (ödenen ÷ ilk muallak)", k["oran"], True),
                                   (f"İlk rezervi yetersiz kalan dosya (ödeme > ilk muallak + %{sapma})", k["yetersiz"], False), (f"İlk rezervi fazla ayrılan dosya (ödeme < ilk muallak − %{sapma})", k["fazla"], False),
                                   ("Ödemesiz kapanan dosya", k["odemesiz"], False), ("Ortalama kapanma süresi (gün)", k["ort_sure"], True)):
        _satir(ws, [etiket, deger], para_s=(2,) if para_mi else ())
    for baslik, alan in (("Açık dosyaların yaşı (gün)", "yaslar"), ("Muallak tutar dilimi", "dilimler")):
        ws.append([])
        _baslik(ws, [baslik, "Dosya", "Muallak Tutar", "Pay %"], dondur="A2")
        for y in r[alan]:
            _satir(ws, [y["dilim"], y["dosya"], y["tutar"], y["pay"]], para_s=(3, 4))
    ws.column_dimensions["C"].width = 18
    ws.append([])
    ws.append([f"Durum tarihi: {bugun:%d.%m.%Y}"])
    ws = wb.create_sheet("Branşlar (Açık)")
    _baslik(ws, ["Branş", "Açık Dosya", "Muallak Tutar", "Pay %", "Ortalama", "En Büyük Dosya", "Ort. Yaş (gün)", "İlk Muallak", "Tahmini Nihai Hasar", "Rezerv Gelişimi %"] + [f"{e} gün" for _, e in YAS], (22, 9, 17, 9, 15, 16, 10, 17, 17, 11, 16, 16, 16, 16), "B2")
    for b in r["branslar"]:
        o = b["acik"]
        if o["dosya"]:
            _satir(ws, [b["brans"], o["dosya"], o["muallak"], o["pay"], o["ort"], o["enb"], o["ort_yas"], o["ilk"], o["nihai"], o["gelisim"]] + [y or None for y in b["yas"]], para_s=range(3, 15),
                   dolgu={10: SARI if o["gelisim"] is not None and o["gelisim"] > sapma else None})
    ws = wb.create_sheet("Rezerv Yeterliliği")
    _baslik(ws, ["Branş", "Kapanan Dosya", "İlk Muallak", "Ödenen", "Gerçekleşme %", "Yetersiz Rezervli Dosya", "Fazla Rezervli Dosya", "Ödemesiz Kapanan", "İlk Rezervi Olmayan", "Ort. Kapanma Süresi (gün)"], (22, 9, 17, 17, 12, 11, 11, 10, 10, 12), "B2")
    for b in r["branslar"]:
        o = b["kapali"]
        if o["dosya"]:
            _satir(ws, [b["brans"], o["dosya"], o["ilk"], o["odenen"], o["oran"], o["yetersiz"] or None, o["fazla"] or None, o["odemesiz"] or None, o["ilksiz"] or None, o["ort_sure"]], para_s=(3, 4, 5, 10),
                   dolgu={5: None if o["oran"] is None else KIRMIZI if o["oran"] > 100 + sapma else SARI if o["oran"] < 100 - sapma else YESIL})
    ws = wb.create_sheet("Açık Dosyalar")
    _baslik(ws, ["Dosya No", "Branş", "İhbar", "Yaş (gün)", "Yaş Dilimi", "İlk Muallak", "Güncel Muallak", "Ödenen", "Tahmini Nihai Hasar", "Tutar Dilimi"], (16, 18, 12, 9, 11, 16, 16, 16, 17, 24))
    for d in r["acik"]:
        _satir(ws, [d["no"], d["brans"], d["ihbar"], d["yas"], d["yas_dilim"], d["ilk"], d["guncel"], d["odenen"] or None, d["nihai"], d["dilim"]], para_s=(6, 7, 8, 9), tarih_s=(3,))
    ws = wb.create_sheet("Kapanan Dosyalar")
    _baslik(ws, ["Dosya No", "Branş", "İhbar", "Kapanış", "Süre (gün)", "İlk Muallak", "Ödenen", "Fark", "Sonuç"], (16, 18, 12, 12, 9, 16, 16, 16, 16))
    for d in r["kapali"]:
        _satir(ws, [d["no"], d["brans"], d["ihbar"], d["kapanis"], d["sure"], d["ilk"], d["odenen"], d["odenen"] - d["ilk"], d["sonuc"]], para_s=(6, 7, 8), tarih_s=(3, 4), dolgu={9: {"Yetersiz": KIRMIZI, "Fazla": SARI, "Uyumlu": YESIL}.get(d["sonuc"])})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, sapma: Decimal = Decimal(20), asgari: int = 3, dilimler: tuple[Decimal, ...] = (Decimal(50000), Decimal(250000), Decimal(1000000))) -> dict:
    dosyalar, uy = oku(yol)
    r = raporla(dosyalar, bugun, sapma, asgari, dilimler)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, bugun, sapma)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Açık hasar dosyalarının branş, yaş ve tutar dağılımını ve kapanan dosyalardan rezerv yeterliliği göstergelerini raporlar.")
    p.add_argument("--dosyalar", type=Path, help="Dosya No, Branş, İhbar Tarihi, Kapanış Tarihi, İlk Muallak, Güncel Muallak, Ödenen Tutar")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--sapma", default="20", help="Rezerv sapma bandı, yüzde (varsayılan 20)")
    p.add_argument("--min", type=int, default=3, dest="asgari", help="Branş uyarısı için en az dosya sayısı (varsayılan 3)")
    p.add_argument("--dilimler", default="50000;250000;1000000", help="Muallak tutar dilim sınırları, artan sırada (varsayılan 50000;250000;1000000)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "muallak_hasar_raporu.xlsx")
    a = p.parse_args(argv)
    ornek = a.dosyalar is None
    yol = ORNEK / "hasar_dosyalari.csv" if ornek else a.dosyalar
    bugun, sapma, dilimler = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today(), para(a.sapma), tuple(para(d) for d in a.dilimler.split(";") if d.strip())
    if bugun is None or sapma is None or not 0 < sapma < 100 or a.asgari < 1 or not dilimler or any(d is None or d <= 0 for d in dilimler) or list(dilimler) != sorted(set(dilimler)):
        print("[X] --tarih GG.AA.YYYY; --sapma 0–100 arası; --min en az 1; --dilimler artan sırada pozitif sayılar olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, bugun, sapma, a.asgari, dilimler)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o, k = r["acik_ozet"], r["kapali_ozet"]
    print(f"[OK] {o['dosya']} açık dosya · muallak {o['muallak']:,.2f}" + (f" · rezerv gelişimi %{o['gelisim']}" if o["gelisim"] is not None else "") + f" · {k['dosya']} kapanan dosya"
          + (f" · gerçekleşme %{k['oran']}" if k["oran"] is not None else ""))
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
