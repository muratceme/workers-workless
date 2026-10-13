import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)


def d(no, yas=10, son=5, tutar=1000, aksiyon="Eksper", sorumlu="S", brans="Kasko"):
    """yas ve son: kaç gün önce (son=None → son işlem yok)."""
    return {"no": no, "brans": brans, "sigortali": "", "ihbar": BUGUN - timedelta(days=yas), "son": None if son is None else BUGUN - timedelta(days=son), "tutar": D(tutar), "aksiyon": aksiyon, "sorumlu": sorumlu}


class DegerlendirTesti(unittest.TestCase):
    def test_oncelik(self):
        l = [d("29", 60, 29), d("30", 60, 30), d("büyük", 60, 30, 250000), d("büyük_altı", 60, 31, 249999), d("sonsuz", 45, None), d("eski179", 179, 1), d("eski180", 180, 1), d("aksiyonsuz", 5, 1, aksiyon=""),
             d("sıfır", 5, 1, 0), d("büyük_hareketli", 5, 1, 900000)]
        r = main.degerlendir(l, BUGUN)
        self.assertEqual({x["no"]: (x["oncelik"], x["bekleme"], x["yas"], x["hareketsiz"]) for x in r["dosyalar"]},
                         {"29": ("Olağan", 29, 60, False), "30": ("Yüksek", 30, 60, True), "büyük": ("Kritik", 30, 60, True), "büyük_altı": ("Yüksek", 31, 60, True), "sonsuz": ("Yüksek", 45, 45, True),
                          "eski179": ("Olağan", 1, 179, False), "eski180": ("Orta", 1, 180, False), "aksiyonsuz": ("Orta", 1, 5, False), "sıfır": ("Orta", 1, 5, False), "büyük_hareketli": ("Olağan", 1, 5, False)})
        self.assertEqual([x["no"] for x in r["dosyalar"]], ["büyük", "sonsuz", "büyük_altı", "30", "aksiyonsuz", "eski180", "sıfır", "29", "büyük_hareketli", "eski179"])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         sorted([("Orta", "Hareketsiz dosya", "30"), ("Yüksek", "Hareketsiz büyük dosya", "büyük"), ("Orta", "Hareketsiz dosya", "büyük_altı"), ("Orta", "Hareketsiz dosya", "sonsuz"), ("Bilgi", "Uzun süredir açık", "eski180"),
                                 ("Orta", "Bekleyen aksiyon yazılmamış", "aksiyonsuz"), ("Orta", "Muallak tutarı sıfır", "sıfır")]))
        p = lambda **kw: {x["no"]: x["oncelik"] for x in main.degerlendir([d("a", 60, 20, 500), d("b", 100, 1)], BUGUN, **kw)["dosyalar"]}   # noqa: E731
        self.assertEqual((p(), p(hareketsiz=20), p(hareketsiz=20, buyuk=D(500)), p(eski=100)), ({"a": "Olağan", "b": "Olağan"}, {"a": "Yüksek", "b": "Olağan"}, {"a": "Kritik", "b": "Olağan"}, {"a": "Olağan", "b": "Orta"}))
        self.assertEqual([u["tur"] for u in main.degerlendir([d("a", 10, -3), d("b", 10, 20)], BUGUN)["uyarilar"]], ["Tarih tutarsız", "Tarih tutarsız"])
        with self.assertRaises(ValueError):
            main.degerlendir([], BUGUN)

    def test_kirilim(self):
        l = [d("1", 30, 10, 100, "Eksper", "A"), d("2", 31, 40, 300, "eksper", "A"), d("3", 90, 2, 50, "Evrak", "B", "Yangın"), d("4", 181, 60, 1000, "", "B", "Yangın"), d("5", 366, 1, 20, "Dava", "B")]
        r = main.degerlendir(l, BUGUN, buyuk=D(1000))
        k = lambda alan: [(s["ad"], s["dosya"], s["tutar"], s["hareketsiz"], s["hareketsiz_tutar"], s["kritik"], s["ort_yas"], s["ort_bekleme"], s["enb_bekleme"]) for s in r[alan]]   # noqa: E731
        self.assertEqual(k("aksiyonlar"), [(main.YOK, 1, D(1000), 1, D(1000), 1, D("181.00"), D("60.00"), 60), ("Eksper", 2, D(400), 1, D(300), 0, D("30.50"), D("25.00"), 40), ("Evrak", 1, D(50), 0, D(0), 0, D("90.00"), D("2.00"), 2),
                                           ("Dava", 1, D(20), 0, D(0), 0, D("366.00"), D("1.00"), 1)])
        self.assertEqual([x[:6] for x in k("sorumlular")], [("B", 3, D(1070), 1, D(1000), 1), ("A", 2, D(400), 1, D(300), 0)])
        self.assertEqual([x[:3] for x in k("branslar")], [("Yangın", 2, D(1050)), ("Kasko", 3, D(420))])
        self.assertEqual([(y["dilim"], y["dosya"], y["tutar"]) for y in r["yaslar"]], [("0–30", 1, D(100)), ("31–90", 2, D(350)), ("91–180", 0, D(0)), ("181–365", 1, D(1000)), ("365+", 1, D(20))])
        self.assertEqual(r["ozet"], {"dosya": 5, "tutar": D(1470), "hareketsiz": 2, "hareketsiz_tutar": D(1300), "kritik": 1, "ort_yas": D("139.60"), "ort_bekleme": D("22.60"), "enb_bekleme": 60})


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "acik_dosyalar.csv", cikti, BUGUN)
            self.assertEqual([(x["no"], x["oncelik"], x["bekleme"], x["yas"]) for x in r["dosyalar"]],
                             [("H-2025-0720", "Kritik", 94, 437), ("H-2026-0042", "Kritik", 41, 216), ("H-2026-0066", "Yüksek", 114, 114), ("H-2026-0087", "Yüksek", 53, 58), ("H-2025-0911", "Orta", 14, 341),
                              ("H-2026-0095", "Orta", 7, 41), ("H-2026-0099", "Olağan", 12, 24), ("H-2026-0101", "Olağan", 3, 10)])
            o = r["ozet"]
            self.assertEqual((o["dosya"], o["tutar"], o["hareketsiz"], o["hareketsiz_tutar"], o["kritik"], o["enb_bekleme"]), (8, D(2060000), 4, D(1440000), 2, 114))
            self.assertEqual([(s["ad"], s["dosya"], s["tutar"]) for s in r["aksiyonlar"]],
                             [("Eksper raporu bekleniyor", 2, D(895000)), ("Dava sonucu bekleniyor", 2, D(710000)), ("Rücu değerlendirmesi", 1, D(275000)), ("Sigortalıdan evrak bekleniyor", 1, D(120000)), (main.YOK, 1, D(60000)),
                              ("Fatura bekleniyor", 1, D(0))])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Yüksek", "Hareketsiz büyük dosya", "H-2025-0720"), ("Yüksek", "Hareketsiz büyük dosya", "H-2026-0042"), ("Orta", "Bekleyen aksiyon yazılmamış", "H-2026-0066"), ("Orta", "Hareketsiz dosya", "H-2026-0066"),
                              ("Orta", "Hareketsiz dosya", "H-2026-0087"), ("Orta", "Muallak tutarı sıfır", "H-2026-0095"), ("Bilgi", "Uzun süredir açık", "H-2025-0911")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Takip Listesi", "Özet", "Bekleyen Aksiyonlar", "Sorumlular", "Branşlar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "d.csv"
            y.write_text("Hasar No;Ürün;Açılış Tarihi;Son Hareket Tarihi;Rezerv;Durum;Hasar Uzmanı\nD1;Kasko;01.09.2026;x;1.500,75;Eksper;Ali\nD2;;02.09.2026;;;;\nd1;;03.09.2026;;;;\nD3;;x;;;;\nD4;;03.09.2026;;-1;;\n;;03.09.2026;;;;\n",
                         encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["brans"], x["ihbar"], x["son"], x["tutar"], x["aksiyon"], x["sorumlu"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("D1", "Kasko", date(2026, 9, 1), None, D("1500.75"), "Eksper", "Ali"), ("D2", "(branş yok)", date(2026, 9, 2), None, D(0), "", "(sorumlu yok)")],
                              [("Tarih okunamadı", "D1"), ("Dosya okunamadı", "d1"), ("Dosya okunamadı", "D3"), ("Dosya okunamadı", "D4")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--dosyalar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "12.10.2026", "--hareketsiz", "45", "--eski", "365", "--buyuk", "500.000"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--hareketsiz", "0"], ["--eski", "0"], ["--buyuk", "0"])], [2] * 4)
            self.assertEqual(main.main(["--dosyalar", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Dosya No;Branş\nA;B\n", "Dosya No;İhbar Tarihi\n;01.03.2026\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
