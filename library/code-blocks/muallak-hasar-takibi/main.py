"""
Muallak Hasar Takibi — Gerisi Bende kod bloğu
Hasar Yönetimi › Hasar Uzmanı

Açık hasar dosyalarını yaş, muallak tutar ve bekleyen aksiyona göre izler; uzun süredir hareketsiz dosyaları öne çıkarır.
  Dosya yaşı (gün)     = durum tarihi − ihbar tarihi;   dilimler 0–30, 31–90, 91–180, 181–365, 365+
  Hareketsiz gün       = durum tarihi − son işlem tarihi (son işlem yoksa ihbar tarihi)
  Hareketsiz dosya     = hareketsiz gün --hareketsiz (30) ve üstünde
  Öncelik              = Kritik (hareketsiz ve muallak ≥ --buyuk, 250.000) · Yüksek (hareketsiz)
                         · Orta (yaş ≥ --eski, 180 gün; aksiyon yazılmamış ya da muallak sıfır) · Olağan
Muallak tutarını ve bekleyen aksiyonu siz yazarsınız; kod rezervin yeterliliğini değerlendirmez.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"no": ("dosya no", "hasar dosya no", "hasar no"), "brans": ("brans", "urun"), "sigortali": ("sigortali", "sigorta ettiren", "musteri"), "ihbar": ("ihbar tarihi", "dosya acilis tarihi", "acilis tarihi"),
         "son": ("son islem tarihi", "son hareket tarihi", "son islem"), "tutar": ("muallak tutar", "muallak", "rezerv", "rezerv tutari"), "aksiyon": ("bekleyen aksiyon", "bekleyen islem", "durum"),
         "sorumlu": ("sorumlu", "dosya sorumlusu", "hasar uzmani")}
YAS = ((30, "0–30"), (90, "31–90"), (180, "91–180"), (365, "181–365"), (None, "365+"))
ONCELIK = ("Kritik", "Yüksek", "Orta", "Olağan")
ONC_RENK = {"Kritik": KIRMIZI, "Yüksek": SARI, "Orta": MAVI, "Olağan": None}
YOK = "(yazılmamış)"


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("no", "ihbar")):
        no = metin(r["no"])
        if not no:
            continue
        ihbar, tutar = tarih(r["ihbar"]), para(r["tutar"]) if metin(r["tutar"]) else Decimal(0)
        if ihbar is None or tutar is None or tutar < 0 or katla(no) in gorulen:
            uy.append(uyari("Yüksek", "Dosya okunamadı", no, f"Satır {r['_satir']}: " + ("aynı dosya numarası daha önce geçti" if katla(no) in gorulen else f"ihbar tarihi '{metin(r['ihbar'])}' ya da muallak tutar '{metin(r['tutar'])}' okunamadı")
                            + ". Satır alınmadı."))
            continue
        gorulen.add(katla(no))
        if metin(r["son"]) and tarih(r["son"]) is None:
            uy.append(uyari("Orta", "Tarih okunamadı", no, f"Son işlem tarihi '{metin(r['son'])}' okunamadı; ihbar tarihi esas alındı."))
        l.append({"no": no, "brans": metin(r["brans"]) or "(branş yok)", "sigortali": metin(r["sigortali"]), "ihbar": ihbar, "son": tarih(r["son"]), "tutar": tutar, "aksiyon": metin(r["aksiyon"]), "sorumlu": metin(r["sorumlu"]) or "(sorumlu yok)"})
    return l, uy


def _ort(d: list[int]) -> Decimal | None:
    return yuvarla(Decimal(sum(d)) / len(d)) if d else None


def degerlendir(dosyalar: list[dict], bugun: date, hareketsiz: int = 30, eski: int = 180, buyuk: Decimal = Decimal(250000)) -> dict:
    if not dosyalar:
        raise ValueError("İzlenecek açık dosya yok.")
    uy = []
    for d in dosyalar:
        d["yas"] = max((bugun - d["ihbar"]).days, 0)
        d["yas_dilim"] = next(e for s, e in YAS if s is None or d["yas"] <= s)
        son = d["son"] or d["ihbar"]
        if son > bugun or son < d["ihbar"]:
            uy.append(uyari("Orta", "Tarih tutarsız", d["no"], f"Son işlem {son:%d.%m.%Y}; ihbar {d['ihbar']:%d.%m.%Y}, durum tarihi {bugun:%d.%m.%Y}. Tarihleri kontrol edin."))
        d["bekleme"] = max((bugun - son).days, 0)
        d["hareketsiz"] = d["bekleme"] >= hareketsiz
        d["oncelik"] = ("Kritik" if d["tutar"] >= buyuk else "Yüksek") if d["hareketsiz"] else "Orta" if d["yas"] >= eski or not d["aksiyon"] or d["tutar"] == 0 else "Olağan"
        bilgi = f"{d['brans']} · muallak {d['tutar']:,.2f} · {d['bekleme']} gündür işlem yok (dosya {d['yas']} günlük). Bekleyen: {d['aksiyon'] or YOK}. Sorumlu: {d['sorumlu']}."
        if d["oncelik"] == "Kritik":
            uy.append(uyari("Yüksek", "Hareketsiz büyük dosya", d["no"], bilgi))
        elif d["hareketsiz"]:
            uy.append(uyari("Orta", "Hareketsiz dosya", d["no"], bilgi))
        elif d["yas"] >= eski:
            uy.append(uyari("Bilgi", "Uzun süredir açık", d["no"], f"Dosya {d['yas']} günlük; muallak {d['tutar']:,.2f}. Bekleyen: {d['aksiyon'] or YOK}."))
        if not d["aksiyon"]:
            uy.append(uyari("Orta", "Bekleyen aksiyon yazılmamış", d["no"], "Dosyanın neyi beklediği belli değil; sıradaki adımı yazın."))
        if d["tutar"] == 0:
            uy.append(uyari("Orta", "Muallak tutarı sıfır", d["no"], "Açık dosyada rezerv yok ya da yazılmamış; tutarı kontrol edin."))

    def ozet(liste: list[dict]) -> dict:
        h = [d for d in liste if d["hareketsiz"]]
        return {"dosya": len(liste), "tutar": sum((d["tutar"] for d in liste), Decimal(0)), "hareketsiz": len(h), "hareketsiz_tutar": sum((d["tutar"] for d in h), Decimal(0)), "kritik": sum(1 for d in liste if d["oncelik"] == "Kritik"),
                "ort_yas": _ort([d["yas"] for d in liste]), "ort_bekleme": _ort([d["bekleme"] for d in liste]), "enb_bekleme": max(d["bekleme"] for d in liste)}

    def kirilim(alan: str) -> list[dict]:
        grup = {}
        for d in dosyalar:
            grup.setdefault(katla(d[alan]), []).append(d)
        return sorted(({"ad": g[0][alan] or YOK, **ozet(g)} for g in grup.values()), key=lambda s: (-s["tutar"], katla(s["ad"])))

    return {"dosyalar": sorted(dosyalar, key=lambda d: (ONCELIK.index(d["oncelik"]), -d["bekleme"], -d["tutar"], katla(d["no"]))), "ozet": ozet(dosyalar), "aksiyonlar": kirilim("aksiyon"), "sorumlular": kirilim("sorumlu"),
            "branslar": kirilim("brans"), "yaslar": [{"dilim": e, "dosya": sum(1 for d in dosyalar if d["yas_dilim"] == e), "tutar": sum((d["tutar"] for d in dosyalar if d["yas_dilim"] == e), Decimal(0))} for _, e in YAS], "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, bugun: date) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Takip Listesi"
    _baslik(ws, ["Öncelik", "Dosya No", "Branş", "Sigortalı", "İhbar", "Son İşlem", "Hareketsiz (gün)", "Yaş (gün)", "Muallak Tutar", "Bekleyen Aksiyon", "Sorumlu", "Not (siz doldurun)"], (10, 16, 14, 24, 12, 12, 11, 9, 16, 36, 20, 36))
    for d in r["dosyalar"]:
        _satir(ws, [d["oncelik"], d["no"], d["brans"], d["sigortali"], d["ihbar"], d["son"], d["bekleme"], d["yas"], d["tutar"], d["aksiyon"] or YOK, d["sorumlu"], None], para_s=(9,), tarih_s=(5, 6),
               dolgu={1: ONC_RENK[d["oncelik"]], 7: SARI if d["hareketsiz"] else None}, sar=(10,))
    ws = wb.create_sheet("Özet")
    o = r["ozet"]
    _baslik(ws, ["Gösterge", "Değer"], (40, 18))
    for etiket, deger, para_mi in (("Açık dosya", o["dosya"], False), ("Toplam muallak", o["tutar"], True), ("Hareketsiz dosya", o["hareketsiz"], False), ("Hareketsiz dosyaların muallağı", o["hareketsiz_tutar"], True),
                                   ("Kritik dosya", o["kritik"], False), ("Ortalama dosya yaşı (gün)", o["ort_yas"], True), ("Ortalama hareketsiz süre (gün)", o["ort_bekleme"], True), ("En uzun hareketsiz süre (gün)", o["enb_bekleme"], False)):
        _satir(ws, [etiket, deger], para_s=(2,) if para_mi else ())
    ws.append([])
    _baslik(ws, ["Dosya yaşı (gün)", "Dosya", "Muallak Tutar"], dondur="A2")
    for y in r["yaslar"]:
        _satir(ws, [y["dilim"], y["dosya"], y["tutar"]], para_s=(3,))
    ws.column_dimensions["C"].width = 18
    ws.append([])
    ws.append([f"Durum tarihi: {bugun:%d.%m.%Y}"])
    for baslik, alan, ad in (("Bekleyen Aksiyonlar", "aksiyonlar", "Bekleyen Aksiyon"), ("Sorumlular", "sorumlular", "Sorumlu"), ("Branşlar", "branslar", "Branş")):
        ws = wb.create_sheet(baslik)
        _baslik(ws, [ad, "Dosya", "Muallak Tutar", "Hareketsiz Dosya", "Hareketsiz Muallak", "Kritik", "Ort. Yaş (gün)", "Ort. Hareketsiz (gün)", "En Uzun Hareketsiz (gün)"], (38, 8, 17, 11, 17, 8, 11, 12, 12))
        for s in r[alan]:
            _satir(ws, [s["ad"], s["dosya"], s["tutar"], s["hareketsiz"] or None, s["hareketsiz_tutar"] or None, s["kritik"] or None, s["ort_yas"], s["ort_bekleme"], s["enb_bekleme"]], para_s=(3, 5, 7, 8), dolgu={6: KIRMIZI if s["kritik"] else None})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, hareketsiz: int = 30, eski: int = 180, buyuk: Decimal = Decimal(250000)) -> dict:
    dosyalar, uy = oku(yol)
    r = degerlendir(dosyalar, bugun, hareketsiz, eski, buyuk)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, bugun)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Açık hasar dosyalarını yaş, muallak tutar ve bekleyen aksiyona göre izler; hareketsiz dosyaları listeler.")
    p.add_argument("--dosyalar", type=Path, help="Dosya No, Branş, Sigortalı, İhbar Tarihi, Son İşlem Tarihi, Muallak Tutar, Bekleyen Aksiyon, Sorumlu")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--hareketsiz", type=int, default=30, help="Kaç gün işlem görmeyen dosya hareketsiz sayılsın (varsayılan 30)")
    p.add_argument("--eski", type=int, default=180, help="Kaç günden eski dosya 'uzun süredir açık' sayılsın (varsayılan 180)")
    p.add_argument("--buyuk", default="250000", help="Büyük dosya sayılacak muallak tutar (varsayılan 250000)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "muallak_hasar_takibi.xlsx")
    a = p.parse_args(argv)
    ornek = a.dosyalar is None
    yol = ORNEK / "acik_dosyalar.csv" if ornek else a.dosyalar
    bugun, buyuk = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today(), para(a.buyuk)
    if bugun is None or a.hareketsiz < 1 or a.eski < 1 or buyuk is None or buyuk <= 0:
        print("[X] --tarih GG.AA.YYYY; --hareketsiz ve --eski en az 1; --buyuk sıfırdan büyük olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, bugun, a.hareketsiz, a.eski, buyuk)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {o['dosya']} açık dosya · muallak {o['tutar']:,.2f} · hareketsiz {o['hareketsiz']} dosya ({o['hareketsiz_tutar']:,.2f}) · kritik {o['kritik']}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
