import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))

import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402


class OrnekTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        with tempfile.TemporaryDirectory() as tmp:
            cls.s = main.calistir(main.ORNEK, Path(tmp) / "r.xlsx", date(2026, 10, 8))
            cls.wb = load_workbook(Path(tmp) / "r.xlsx")
        cls.y = {y.no[-4:]: y for y in cls.s["yazilar"]}

    def notlar(self, no):
        return "\n".join(f"{a} {b}" for a, b in self.y[no].notlar)

    def test_turler(self):
        self.assertEqual([self.y[n].tur for n in ("0101", "0105", "0107", "0108", "0110", "0111")],
                         ["İİK 89/1 haciz ihbarnamesi", "İİK 89/2 haciz ihbarnamesi", "6183 haciz bildirisi", "İİK 89/3 haciz ihbarnamesi",
                          "Haciz kaldırma", "Bilgi / belge talebi"])

    def test_son_gunler(self):
        self.assertEqual(self.y["0102"].son_gun, date(2026, 10, 2))       # 25.09 + 7 gün
        self.assertEqual(self.y["0103"].son_gun, date(2026, 10, 5))       # 26.09 + 7 = 03.10 Cumartesi → Pazartesi
        self.assertEqual(self.y["0105"].son_gun, date(2026, 10, 19))      # 02.10 + 15 = 17.10 Cumartesi → 19.10
        self.assertEqual(self.y["0107"].son_gun, date(2026, 10, 12))      # yazıda 7 gün
        self.assertEqual(self.y["0112"].son_gun, date(2026, 10, 8))       # fekk: 1 iş günü
        self.assertIsNone(self.y["0109"].son_gun)

    def test_durumlar(self):
        d = {n: y.durum for n, y in self.y.items()}
        self.assertEqual(d["0101"], "Cevaplandı")
        self.assertEqual(d["0102"], "Süresi geçti")
        self.assertEqual(d["0103"], "Geç cevaplandı")
        self.assertEqual(d["0104"], "Bugün son gün")
        self.assertEqual(d["0106"], "Süresi geçti")
        self.assertEqual((d["0107"], self.y["0107"].kalan), ("Yaklaşan", 2))
        self.assertEqual(d["0109"], "Süre girilmeli")

    def test_zincir_ve_kontroller(self):
        self.assertIn("Kritik 89/1 ihbarnamesi (RY-2026-0102) süresinde cevaplanmamış", self.notlar("0105"))
        self.assertIn("Kritik Üçüncü haciz ihbarnamesi", self.notlar("0108"))
        self.assertIn("Hata TCKN kontrol haneleri tutmuyor", self.notlar("0103"))
        self.assertIn("Mükerrer kayıt olabilir: RY-2026-0104", self.notlar("0113"))
        self.assertIn("haciz kaldırma yazısı da var", self.notlar("0104"))

    def test_musteri_ozeti(self):
        m = {x["musteri"]: x for x in self.s["ozet"]}
        self.assertEqual((m["M200"]["acik"], m["M200"]["haciz_tutar"]), (3, D("1590000.00")))   # 89/2 tekrar sayılmaz
        self.assertEqual(m["M300"]["haciz_tutar"], D("27300.00"))                              # mükerrer sayılmaz
        self.assertEqual(self.s["ozet"][0]["musteri"], "M200")

    def test_excel(self):
        self.assertEqual(self.wb.sheetnames, ["Özet", "Takip Listesi", "Müşteri Özeti", "Süre Kuralları"])
        self.assertEqual(self.wb["Takip Listesi"]["A2"].value, "Süresi geçti")


class TakvimTesti(unittest.TestCase):
    def test_tatil(self):
        t = main.tatil_kumesi({2026}, set())
        self.assertEqual(main.ilk_is_gunu(date(2026, 10, 29), t), date(2026, 10, 30))          # Cumhuriyet Bayramı
        self.assertEqual(main.ilk_is_gunu(date(2026, 5, 27), t), date(2026, 6, 1))             # Kurban Bayramı + hafta sonu
        self.assertEqual(main.is_gunu_ekle(date(2026, 10, 28), 1, t), date(2026, 10, 30))
        self.assertEqual(main.is_gunu_farki(date(2026, 10, 8), date(2026, 10, 12), t), 2)
        self.assertEqual(main.is_gunu_farki(date(2026, 10, 12), date(2026, 10, 8), t), -2)

    def test_kimlik(self):
        self.assertEqual(main.kimlik_kontrol("10000000146"), "")
        self.assertEqual(main.kimlik_kontrol("1234567890"), "")
        self.assertIn("VKN", main.kimlik_kontrol("1234567891"))
        self.assertIn("hane", main.kimlik_kontrol("123"))


if __name__ == "__main__":
    unittest.main()
