"""
Resmî Yazı (Haciz/Müzekkere) Takibi — Workers / Workless kod bloğu
Bankacılık › Şube Bankacılığı › Operasyon Yetkilisi

İcra dairesi, vergi dairesi, SGK, mahkeme ve savcılıktan gelen haciz ihbarnamesi, haciz bildirisi, haciz
kaldırma ve bilgi talebi yazılarını cevap süresiyle takip eder:
  - Yazı türü metinden tanınır (İİK 89/1, 89/2, 89/3, 6183 haciz bildirisi, haciz kaldırma, bilgi talebi, tedbir).
  - Son cevap günü: yazıda süre yazıyorsa o, yoksa yasal süre (İİK 89/1: 7 gün, 89/2: 15 gün; tebliğden itibaren
    takvim günü). Son gün hafta sonuna veya resmî tatile denk gelirse izleyen ilk iş günü. Haciz kaldırma için
    iç hedef (--fekk-gun iş günü). Süresi bilinmeyen yazılar "Süre girilmeli" olarak işaretlenir.
  - Durum: süresi geçti, bugün son gün, yaklaşan (--uyari-gun), süresi var, cevaplandı / geç cevaplandı.
  - Zincir kontrolü: aynı dosyada 89/1 cevapsızken 89/2 gelmesi, 89/3 gelmesi (kritik); haciz kaldırma gelen
    dosyalar; mükerrer kayıt.
  - TCKN/VKN kontrol hanesi; müşteri bazında açık yazı ve haciz tutarı özeti.
İnternete bağlanmaz.

Kullanım:
    python main.py                                                     # örnek yazılar (rapor tarihi 08.10.2026)
    python main.py --girdi gelen_yazilar.xlsx --tarih 09.10.2026 --tatiller tatiller.csv
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri" / "gelen_yazilar.csv"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")

# (tür, anahtar ifadeler — katlanmış metinde, yasal süre (takvim günü) veya None, dayanak/açıklama)
TURLER = [
    ("İİK 89/3 haciz ihbarnamesi", ("89/3", "ucuncu haciz ihbar", "3. haciz ihbar"), None,
     "Üçüncü ihbarname: zimmet/sorumluluk riski. Hukuk birimine iletin; süre yazıdan girilmeli."),
    ("İİK 89/2 haciz ihbarnamesi", ("89/2", "ikinci haciz ihbar", "2. haciz ihbar"), 15,
     "İİK 89/2: tebliğden itibaren 15 gün içinde menfi tespit davası açılabilir."),
    ("İİK 89/1 haciz ihbarnamesi", ("89/1", "birinci haciz ihbar", "1. haciz ihbar", "haciz ihbarname"), 7,
     "İİK 89/1: tebliğden itibaren 7 gün içinde itiraz/beyan."),
    ("Haciz kaldırma", ("kaldir", "fekk"), None, "İç hedef: --fekk-gun iş günü içinde işlenmeli."),
    ("6183 haciz bildirisi", ("6183", "haciz bildiri"), None, "Amme alacağı (vergi dairesi, SGK): süre yazıdan girilmeli."),
    ("Tedbir kararı", ("tedbir",), None, "Süre ve kapsam yazıdan girilmeli."),
    ("Bilgi / belge talebi", ("bilgi", "belge talep", "muzekkere"), None, "Süre yazıdan girilmeli."),
]
HACIZ_TURLERI = {"İİK 89/1 haciz ihbarnamesi", "İİK 89/2 haciz ihbarnamesi", "İİK 89/3 haciz ihbarnamesi", "6183 haciz bildirisi"}

# Genel tatiller (2912 sayılı Kanun) ve dini bayramlar (Diyanet takvimi; her yıl --tatiller ile kontrol edin)
SABIT_TATILLER = {(1, 1), (4, 23), (5, 1), (5, 19), (7, 15), (8, 30), (10, 29)}
DINI_BAYRAMLAR = {2025: [(date(2025, 3, 30), 3), (date(2025, 6, 6), 4)], 2026: [(date(2026, 3, 20), 3), (date(2026, 5, 27), 4)],
                  2027: [(date(2027, 3, 9), 3), (date(2027, 5, 16), 4)]}

SUTUNLAR = {
    "no": ("kayit no", "evrak no", "no", "sira no"),
    "teblig": ("teblig tarihi", "gelis tarihi", "tarih", "evrak tarihi", "giris tarihi"),
    "kurum": ("kurum", "gonderen kurum", "gonderen", "mercii"),
    "tur": ("yazi turu", "tur", "evrak turu", "konu"),
    "dosya": ("dosya no", "esas no", "dosya"),
    "musteri": ("musteri no", "musteri"),
    "ad": ("musteri adi", "borclu", "borclu adi", "unvan", "ad soyad"),
    "kimlik": ("tckn/vkn", "tckn", "vkn", "kimlik no", "tc kimlik no", "vergi no"),
    "tutar": ("tutar", "haciz tutari", "alacak tutari"),
    "sure": ("yazidaki sure (gun)", "yazidaki sure", "sure (gun)", "sure"),
    "cevap": ("cevap tarihi", "yanit tarihi", "islem tarihi"),
    "sorumlu": ("sorumlu", "sorumlu birim"),
}


def katla(s) -> str:
    return " ".join(str(s or "").translate(_TR).lower().split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%Y %H:%M"):
        try:
            return datetime.strptime(str(x or "").strip(), f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif s.count(".") > 1 or (s.count(".") == 1 and len(s.split(".")[1]) == 3):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kimlik_kontrol(no: str) -> str:
    """Boş metin = geçerli (veya boş). TCKN 11 hane, VKN 10 hane kontrol hanesi."""
    if not no:
        return ""
    if re.fullmatch(r"\d{11}", no):
        d = [int(c) for c in no]
        if d[0] == 0 or d[9] != (sum(d[0:9:2]) * 7 - sum(d[1:8:2])) % 10 or d[10] != sum(d[:10]) % 10:
            return "TCKN kontrol haneleri tutmuyor"
        return ""
    if re.fullmatch(r"\d{10}", no):
        s = 0
        for i, n in enumerate(reversed(no[:9]), 1):
            c1 = (int(n) + i) % 10
            if c1:
                s += (c1 * 2 ** i) % 9 or 9
        return "" if (10 - s % 10) % 10 == int(no[9]) else "VKN kontrol hanesi tutmuyor"
    return f"TCKN 11, VKN 10 hane olmalı ({len(no)} hane)"


# ----------------------------------------------------------------------------
# Takvim
# ----------------------------------------------------------------------------

def tatil_kumesi(yillar: set[int], ek: set[date]) -> set[date]:
    t = set(ek)
    for y in yillar:
        t |= {date(y, a, g) for a, g in SABIT_TATILLER}
        for ilk, n in DINI_BAYRAMLAR.get(y, []):
            t |= {ilk + timedelta(days=i) for i in range(n)}
    return t


def is_gunu_mu(d: date, tatiller: set[date]) -> bool:
    return d.weekday() < 5 and d not in tatiller


def ilk_is_gunu(d: date, tatiller: set[date]) -> date:
    while not is_gunu_mu(d, tatiller):
        d += timedelta(days=1)
    return d


def is_gunu_ekle(d: date, n: int, tatiller: set[date]) -> date:
    while n > 0:
        d += timedelta(days=1)
        if is_gunu_mu(d, tatiller):
            n -= 1
    return d


def is_gunu_farki(bas: date, son: date, tatiller: set[date]) -> int:
    """bas'tan sonraki günden son'a kadar (son dahil) iş günü sayısı; son geçmişse eksi değer."""
    if son < bas:
        return -is_gunu_farki(son, bas, tatiller)
    return sum(1 for i in range(1, (son - bas).days + 1) if is_gunu_mu(bas + timedelta(days=i), tatiller))


# ----------------------------------------------------------------------------
# Yazılar
# ----------------------------------------------------------------------------

@dataclass
class Yazi:
    no: str
    teblig: date | None
    kurum: str
    tur_metni: str
    dosya: str
    musteri: str
    ad: str
    kimlik: str
    tutar: Decimal | None
    yazi_suresi: int | None
    cevap: date | None
    sorumlu: str
    tur: str = "Diğer"
    dayanak: str = ""
    son_gun: date | None = None
    sure_kaynagi: str = ""
    durum: str = ""
    kalan: int | None = None
    notlar: list[tuple[str, str]] = field(default_factory=list)      # (önem, açıklama)

    @property
    def onem(self) -> str:
        sira = ["Kritik", "Hata", "Yüksek", "Dikkat", "Bilgi"]
        return min((n[0] for n in self.notlar), key=sira.index, default="")


def tur_bul(metin: str) -> tuple[str, int | None, str]:
    k = katla(metin)
    for ad, anahtarlar, sure, dayanak in TURLER:
        if any(a in k for a in anahtarlar):
            return ad, sure, dayanak
    return "Diğer", None, "Süre yazıdan girilmeli."


def yazilari_oku(yol: Path) -> list[Yazi]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:10]):
        b = [katla(x) for x in r]
        k = {a: next((i for i, x in enumerate(b) if x in es), None) for a, es in SUTUNLAR.items()}
        if k["teblig"] is not None and k["tur"] is not None:
            break
    else:
        raise SystemExit(f"{yol.name}: 'Tebliğ Tarihi' ve 'Yazı Türü' sütunları bulunamadı. İlk satır: {s[0] if s else '(boş)'}")
    al = lambda r, a: r[k[a]] if k[a] is not None and k[a] < len(r) else None  # noqa: E731
    yazilar = []
    for i, r in enumerate(s[bi + 1:], 1):
        sure = para(al(r, "sure"))
        y = Yazi(str(al(r, "no") or f"S{i}").strip(), tarih(al(r, "teblig")), str(al(r, "kurum") or "").strip(), str(al(r, "tur") or "").strip(),
                 str(al(r, "dosya") or "").strip(), str(al(r, "musteri") or "").strip(), str(al(r, "ad") or "").strip(),
                 re.sub(r"\D", "", str(al(r, "kimlik") or "")), para(al(r, "tutar")), int(sure) if sure else None, tarih(al(r, "cevap")),
                 str(al(r, "sorumlu") or "").strip())
        y.tur, _, y.dayanak = tur_bul(y.tur_metni)
        yazilar.append(y)
    return yazilar


def degerlendir(yazilar: list[Yazi], bugun: date, tatiller: set[date], uyari_gun: int, fekk_gun: int) -> None:
    for y in yazilar:
        _, yasal, _ = tur_bul(y.tur_metni)
        if y.teblig is None:
            y.notlar.append(("Hata", "Tebliğ tarihi yok veya okunamadı"))
        elif y.teblig > bugun:
            y.notlar.append(("Hata", f"Tebliğ tarihi ({y.teblig:%d.%m.%Y}) rapor tarihinden sonra"))
        h = kimlik_kontrol(y.kimlik)
        if h:
            y.notlar.append(("Hata", h))
        if y.teblig:
            if y.yazi_suresi:
                y.son_gun, y.sure_kaynagi = ilk_is_gunu(y.teblig + timedelta(days=y.yazi_suresi), tatiller), f"Yazıda {y.yazi_suresi} gün"
                if yasal and y.yazi_suresi != yasal:
                    y.notlar.append(("Bilgi", f"Yazıdaki süre ({y.yazi_suresi} gün) yasal süreden ({yasal} gün) farklı; yazıdaki esas alındı"))
            elif yasal:
                y.son_gun, y.sure_kaynagi = ilk_is_gunu(y.teblig + timedelta(days=yasal), tatiller), f"Yasal {yasal} gün"
            elif y.tur == "Haciz kaldırma":
                y.son_gun, y.sure_kaynagi = is_gunu_ekle(y.teblig, fekk_gun, tatiller), f"İç hedef {fekk_gun} iş günü"
            if y.son_gun and y.teblig + timedelta(days=y.yazi_suresi or yasal or 0) != y.son_gun and y.sure_kaynagi.startswith(("Yazıda", "Yasal")):
                y.notlar.append(("Bilgi", f"Son gün tatile denk geldiği için {y.son_gun:%d.%m.%Y} iş gününe uzadı"))
        if y.cevap:
            if y.teblig and y.cevap < y.teblig:
                y.notlar.append(("Hata", "Cevap tarihi tebliğden önce"))
            y.durum = "Geç cevaplandı" if y.son_gun and y.cevap > y.son_gun else "Cevaplandı"
            if y.durum == "Geç cevaplandı":
                y.notlar.append(("Yüksek", f"Son günden ({y.son_gun:%d.%m.%Y}) {(y.cevap - y.son_gun).days} gün sonra cevaplanmış"))
        elif y.son_gun is None:
            y.durum = "Süre girilmeli"
            y.notlar.append(("Dikkat", "Yasal veya yazıda belirtilen süre yok; yazıdaki süreyi girin"))
        else:
            y.kalan = is_gunu_farki(bugun, y.son_gun, tatiller)
            if y.son_gun < bugun:
                y.durum = "Süresi geçti"
                y.notlar.append(("Kritik", f"Son gün {y.son_gun:%d.%m.%Y} geçti, cevap kaydı yok"))
            elif y.son_gun == bugun:
                y.durum = "Bugün son gün"
                y.notlar.append(("Yüksek", "Bugün son gün"))
            elif y.kalan <= uyari_gun:
                y.durum = "Yaklaşan"
                y.notlar.append(("Dikkat", f"{y.kalan} iş günü kaldı"))
            else:
                y.durum = "Süresi var"
        if y.tur == "İİK 89/3 haciz ihbarnamesi":
            y.notlar.append(("Kritik", "Üçüncü haciz ihbarnamesi: hukuk birimine iletin"))
    # Zincir ve mükerrer
    anahtar = lambda y: (katla(y.dosya), y.kimlik or katla(y.ad))  # noqa: E731
    dosya = defaultdict(list)
    for y in yazilar:
        if y.dosya:
            dosya[anahtar(y)].append(y)
    for ys in dosya.values():
        birinci = [y for y in ys if y.tur == "İİK 89/1 haciz ihbarnamesi"]
        for y in ys:
            if y.tur == "İİK 89/2 haciz ihbarnamesi":
                if not birinci:
                    y.notlar.append(("Bilgi", "Bu dosyanın 89/1 kaydı listede yok"))
                elif all(b.cevap is None or (b.son_gun and b.cevap > b.son_gun) for b in birinci):
                    y.notlar.append(("Kritik", f"89/1 ihbarnamesi ({birinci[0].no}) süresinde cevaplanmamış görünüyor; 89/2 geldi — "
                                               "menfi tespit süresi için hukuk birimine iletin"))
        if any(y.tur == "Haciz kaldırma" for y in ys):
            for y in ys:
                if y.tur in HACIZ_TURLERI and y.cevap is None:
                    y.notlar.append(("Bilgi", "Bu dosyada haciz kaldırma yazısı da var; işlem sırasını kontrol edin"))
        gorulen = {}
        for y in ys:
            k = (y.tur, y.teblig)
            if k in gorulen:
                y.notlar.append(("Dikkat", f"Mükerrer kayıt olabilir: {gorulen[k]} ile aynı dosya, tür ve tebliğ tarihi"))
            else:
                gorulen[k] = y.no


def musteri_ozeti(yazilar: list[Yazi]) -> list[dict]:
    m = defaultdict(lambda: {"ad": "", "kimlik": "", "yazi": 0, "acik": 0, "haciz_tutar": Decimal(0), "kurumlar": set(), "en_yakin": None})
    for y in yazilar:
        x = m[y.musteri or y.kimlik or y.ad]
        x["ad"], x["kimlik"] = y.ad, y.kimlik
        x["yazi"] += 1
        x["kurumlar"].add(y.kurum)
        if y.cevap is None:
            x["acik"] += 1
            if y.son_gun and (x["en_yakin"] is None or y.son_gun < x["en_yakin"]):
                x["en_yakin"] = y.son_gun
        if y.tur in HACIZ_TURLERI and y.tutar and not any("Mükerrer" in n[1] for n in y.notlar) and y.tur != "İİK 89/2 haciz ihbarnamesi":
            x["haciz_tutar"] += y.tutar                       # 89/2 aynı alacak için tekrar sayılmaz
    return sorted(({"musteri": k, **v} for k, v in m.items()), key=lambda x: (-x["acik"], -x["haciz_tutar"]))


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Kritik": "F8C9C6", "Hata": "FDE2E1", "Yüksek": "FDE2E1", "Dikkat": "FFF4CE", "Bilgi": "E8F0FE",
        "Süresi geçti": "F8C9C6", "Bugün son gün": "FDE2E1", "Yaklaşan": "FFF4CE", "Süre girilmeli": "FFF4CE", "Geç cevaplandı": "FDE2E1",
        "Cevaplandı": "E3F5E1", "Süresi var": "FFFFFF"}
UST = Alignment(vertical="top", wrap_text=True)
DURUM_SIRA = ["Süresi geçti", "Bugün son gün", "Yaklaşan", "Süre girilmeli", "Süresi var", "Geç cevaplandı", "Cevaplandı"]


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def rapor_yaz(cikti: Path, yazilar: list[Yazi], ozet: list[dict], bugun: date) -> None:
    wb = Workbook()
    o = wb.active
    o.title = "Özet"
    o.column_dimensions["A"].width, o.column_dimensions["B"].width = 30, 12
    o.append(["Rapor tarihi", bugun])
    o["B1"].number_format = "DD.MM.YYYY"
    o.append([])
    _baslik(o, ["Durum", "Yazı"], ())
    say = Counter(y.durum for y in yazilar)
    for d in DURUM_SIRA:
        if say[d]:
            o.append([d, say[d]])
            o.cell(o.max_row, 1).fill = PatternFill("solid", fgColor=RENK[d])
    o.append([])
    _baslik(o, ["Yazı türü", "Yazı", "Açık"], ())
    tur = Counter(y.tur for y in yazilar)
    for t, n in tur.most_common():
        o.append([t, n, sum(1 for y in yazilar if y.tur == t and y.cevap is None)])
    o.append([])
    o.append(["Not", "Süreler örnektir: yazıdaki süre ve hukuk biriminizin görüşü esastır. Yasal süreler tebliğden itibaren takvim günüyle "
                     "sayılır; son gün tatile denk gelirse izleyen iş gününe uzar. Dini bayramları her yıl --tatiller ile kontrol edin."])
    o.cell(o.max_row, 2).alignment = UST
    o.column_dimensions["B"].width = 80

    t = wb.create_sheet("Takip Listesi")
    _baslik(t, ["Durum", "Son Gün", "Kalan İş Günü", "Kayıt No", "Tebliğ", "Kurum", "Tür", "Dosya No", "Müşteri No", "Müşteri", "TCKN/VKN",
                "Tutar", "Süre Kaynağı", "Notlar", "Sorumlu", "Cevap Tarihi", "Yapılan İşlem"],
            (14, 11, 9, 14, 11, 28, 26, 18, 10, 24, 13, 14, 16, 60, 13, 12, 26))
    for y in sorted(yazilar, key=lambda y: (DURUM_SIRA.index(y.durum), y.son_gun or date.max)):
        t.append([y.durum, y.son_gun, y.kalan, y.no, y.teblig, y.kurum, y.tur, y.dosya, y.musteri, y.ad, y.kimlik,
                  None if y.tutar is None else float(y.tutar), y.sure_kaynagi, "\n".join(f"[{a}] {b}" for a, b in y.notlar), y.sorumlu, y.cevap, ""])
        t.cell(t.max_row, 1).fill = PatternFill("solid", fgColor=RENK[y.durum])
        for c in (2, 5, 16):
            t.cell(t.max_row, c).number_format = "DD.MM.YYYY"
        t.cell(t.max_row, 12).number_format = "#,##0.00"
        t.cell(t.max_row, 11).number_format = "@"
        t.cell(t.max_row, 17).fill = PatternFill("solid", fgColor="FFF4CE")
        if y.onem:
            t.cell(t.max_row, 14).fill = PatternFill("solid", fgColor=RENK[y.onem])
        for h in t[t.max_row]:
            h.alignment = UST
    t.freeze_panes = "E2"
    t.auto_filter.ref = t.dimensions

    m = wb.create_sheet("Müşteri Özeti")
    _baslik(m, ["Müşteri No", "Müşteri", "TCKN/VKN", "Yazı", "Açık", "Haciz Tutarı (89/1, 89/3, 6183)", "Kurum Sayısı", "En Yakın Son Gün"],
            (11, 26, 13, 7, 7, 22, 12, 15))
    for x in ozet:
        m.append([x["musteri"], x["ad"], x["kimlik"], x["yazi"], x["acik"], float(x["haciz_tutar"]) or None, len(x["kurumlar"]), x["en_yakin"]])
        m.cell(m.max_row, 6).number_format = "#,##0.00"
        m.cell(m.max_row, 8).number_format = "DD.MM.YYYY"
        if len(x["kurumlar"]) >= 2 and x["haciz_tutar"]:
            m.cell(m.max_row, 7).fill = PatternFill("solid", fgColor=RENK["Dikkat"])

    k = wb.create_sheet("Süre Kuralları")
    _baslik(k, ["Tür", "Yasal süre (takvim günü)", "Dayanak / not"], (28, 22, 90))
    for ad, _, sure, dayanak in TURLER:
        k.append([ad, sure or "yazıdan", dayanak])
    k.append(["Diğer", "yazıdan", "Tanınmayan yazı türü"])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(girdi: Path, cikti: Path, bugun: date | None = None, ek_tatiller: set[date] | None = None, uyari_gun: int = 2,
             fekk_gun: int = 1) -> dict:
    bugun = bugun or date.today()
    yazilar = yazilari_oku(girdi)
    yillar = {bugun.year, bugun.year + 1} | {y.teblig.year for y in yazilar if y.teblig}
    tatiller = tatil_kumesi(yillar, ek_tatiller or set())
    degerlendir(yazilar, bugun, tatiller, uyari_gun, fekk_gun)
    ozet = musteri_ozeti(yazilar)
    rapor_yaz(cikti, yazilar, ozet, bugun)
    return {"yazilar": yazilar, "ozet": ozet, "bugun": bugun}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Haciz ihbarnamesi, haciz bildirisi ve müzekkereleri cevap süresiyle takip eder.")
    p.add_argument("--girdi", type=Path, default=ORNEK, help="Gelen yazılar (.xlsx/.csv)")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY (varsayılan: bugün; örnek veride 08.10.2026)")
    p.add_argument("--tatiller", type=Path, help="Ek tatil günleri (.csv/.txt, satır başına GG.AA.YYYY)")
    p.add_argument("--uyari-gun", type=int, default=2, help="Son güne bu kadar iş günü kala 'Yaklaşan' (varsayılan 2)")
    p.add_argument("--fekk-gun", type=int, default=1, help="Haciz kaldırma yazısı için iç hedef, iş günü (varsayılan 1)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "resmi_yazi_takibi.xlsx")
    a = p.parse_args(argv)
    bugun = tarih(a.tarih) if a.tarih else (date(2026, 10, 8) if a.girdi == ORNEK else None)
    if a.tarih and bugun is None:
        print(f"[X] Tarih GG.AA.YYYY olmalı: {a.tarih}")
        return 2
    ek = set()
    if a.tatiller:
        ek = {d for d in (tarih(x) for x in re.findall(r"\d{1,2}\.\d{1,2}\.\d{4}|\d{4}-\d\d-\d\d", a.tatiller.read_text(encoding="utf-8-sig"))) if d}
    s = calistir(a.girdi, a.cikti, bugun, ek, a.uyari_gun, a.fekk_gun)
    d = Counter(y.durum for y in s["yazilar"])
    print(f"[OK] {len(s['yazilar'])} yazı · rapor tarihi {s['bugun']:%d.%m.%Y} · " + " · ".join(f"{k} {d[k]}" for k in DURUM_SIRA if d[k]))
    for y in s["yazilar"]:
        for o, n in y.notlar:
            if o in ("Kritik", "Hata"):
                print(f"[X] {y.no} ({y.tur}, {y.musteri}): {n}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
