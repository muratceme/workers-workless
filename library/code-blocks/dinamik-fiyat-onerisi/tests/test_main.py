import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
B = date(2026, 10, 12)


def s(gun, satilan, pickup, mevcut, rakip=None, taban=None, tavan=None, kapasite=100, oda="Std", ay=10):
    f = lambda v: None if v is None else D(v)   # noqa: E731
    return {"tarih": date(2026, ay, gun), "oda": oda, "kapasite": D(kapasite), "satilan": D(satilan), "pickup": f(pickup), "mevcut": D(mevcut), "rakip": f(rakip), "taban": f(taban), "tavan": f(tavan)}


def tek(*a, **kw):
    ayar = {k: kw.pop(k) for k in list(kw) if k in main.AYAR}
    return main.oner([s(*a, **kw)], B, **ayar)["oneriler"][0]


class OneriTesti(unittest.TestCase):
    def test_temel_artis(self):
        x = tek(19, 80, 14, 3000, 3200)
        self.assertEqual((x["kalan"], x["doluluk"], x["hiz"], x["tahmin"], x["ayar"], x["oneri"], x["degisim"], x["tur"], x["rakip_fark"]), (7, D("80.00"), D("2.00"), D("94.00"), D("4.50"), D("3140.00"), D("4.67"), "Artır", D("-1.88")))
        self.assertEqual(x["gerekce"], "Tahmini doluluk %94.00, hedef %85: ayar +4.50%.")
        self.assertEqual((tek(19, 80, 14, 3000, yuvarla=D(0))["oneri"], tek(19, 80, 14, 3000, duyarlilik=D(1), azami=D(5))["oneri"], tek(19, 80, 14, 3000, hedef=D(94))["tur"], tek(19, 80, 14, 3000, pickup_gun=14)["tahmin"],
                          tek(19, 80, 14, 3000, projeksiyon=2)["tahmin"]), (D("3135.00"), D("3150.00"), "Koru", D("87.00"), D("84.00")))

    def test_indirim_ve_sinirlar(self):
        x = tek(26, 30, 7, 3000, 3200, 2800)
        self.assertEqual((x["kalan"], x["tahmin"], x["ayar"], x["oneri"], x["degisim"], x["tur"]), (14, D("44.00"), D("-20.00"), D("2800.00"), D("-6.67"), "İndir"))
        self.assertIn("azami sınırda", x["gerekce"])
        self.assertIn("Rakip bandının alt sınırı 2,720.00", x["gerekce"])
        self.assertIn("Taban fiyat 2,800.00", x["gerekce"])
        self.assertEqual((tek(26, 30, 7, 3000, 3200)["oneri"], tek(26, 30, 7, 3000)["oneri"], tek(26, 30, 7, 3000, 3205)["oneri"], tek(26, 30, 7, 2500, 3200)["oneri"], tek(26, 30, 7, 2500, 3200)["tur"]),
                         (D("2720.00"), D("2400.00"), D("2730.00"), D("2500.00"), "Koru"))                # 3205 × 0,85 = 2.724,25 → yukarı 2.730; mevcut zaten bandın altındaysa indirilmez
        x = tek(22, 70, -14, 2000)
        self.assertEqual((x["hiz"], x["tahmin"], x["ayar"], x["oneri"], x["degisim"]), (D("-2.00"), D("50.00"), D("-17.50"), D("1650.00"), D("-17.50")))
        self.assertEqual(tek(22, 5, -70, 2000)["tahmin"], D("0.00"))

    def test_rakip_ust_bandi_tavan_ve_dolu(self):
        x = tek(11, 60, 21, 3000, 2700, ay=11)
        self.assertEqual((x["kalan"], x["tahmin"], x["ayar"], x["oneri"], x["degisim"], x["tur"]), (30, D("102.00"), D("8.50"), D("3100.00"), D("3.33"), "Artır"))
        r = main.oner([s(20, 90, 7, 4000, 3000)], B)
        x = r["oneriler"][0]
        self.assertEqual((x["tahmin"], x["oneri"], x["tur"], [u["tur"] for u in r["uyarilar"]]), (D("98.00"), D("4000.00"), "Koru", ["Mevcut fiyat rakip bandının üstünde"]))
        x = tek(13, 10, 7, 5000, tavan=5200, kapasite=10)
        self.assertEqual((x["tahmin"], x["ayar"], x["oneri"], x["degisim"]), (D("110.00"), D("12.50"), D("5200.00"), D("4.00")))
        self.assertIn("Kapasite dolu", x["gerekce"])
        self.assertEqual((tek(13, 85, None, 2500)["tur"], tek(13, 85, None, 2500)["oneri"], tek(12, 100, 70, 1000)["tahmin"], tek(13, 85, 0, 2500, taban=2600)["oneri"]), ("Koru", D("2500.00"), D("100.00"), D("2600.00")))

    def test_siralama_gecmis_ve_hatalar(self):
        r = main.oner([s(20, 110, 0, 1000), s(11, 50, 0, 1000), s(19, 50, 0, 1000, oda="Suit"), s(19, 50, 0, 1000, oda="Aile")], B)
        self.assertEqual(([(x["tarih"].day, x["oda"]) for x in r["oneriler"]], r["say"], sorted(u["tur"] for u in r["uyarilar"])),
                         ([(19, "Aile"), (19, "Suit"), (20, "Std")], {"Artır": 1, "İndir": 2, "Koru": 0}, ["Geçmiş tarih", "Satılan oda kapasiteden fazla"]))
        for l in ([], [s(11, 50, 0, 1000)]):
            with self.assertRaises(ValueError):
                main.oner(l, B)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "f.xlsx"
            r = main.calistir(ORNEK / "durum.csv", cikti, B)
            self.assertEqual([(f"{x['tarih']:%d.%m}", x["oda"], x["oneri"], x["tur"]) for x in r["oneriler"]],
                             [("13.10", "Standart", D("2500.00"), "Koru"), ("19.10", "Standart", D("3140.00"), "Artır"), ("19.10", "Suit", D("5200.00"), "Artır"), ("20.10", "Standart", D("4000.00"), "Koru"),
                              ("26.10", "Standart", D("2800.00"), "İndir"), ("11.11", "Standart", D("3100.00"), "Artır")])
            self.assertEqual((r["say"], [(u["tur"], u["kim"]) for u in r["uyarilar"]]), ({"Artır": 3, "İndir": 1, "Koru": 2}, [("Mevcut fiyat rakip bandının üstünde", "20.10.2026 · Standart")]))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Fiyat Önerisi", "Kurallar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "d.csv"
            y.write_text("Tarih;Kapasite;Satılan;Pick-up;Fiyat;Rakip;Taban;Tavan\n20.10.2026;100;50;-3;1.250,50;;;\n21.10.2026;0;5;1;100;;;\n22.10.2026;10;5;x;100;;;\n23.10.2026;10;5;1;100;;200;150\nyarın;10;5;1;100;;;\n"
                         "24.10.2026;10;;1;100;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["tarih"].day, x["oda"], x["pickup"], x["mevcut"], x["rakip"]) for x in l], [u["tur"] for u in uy]), ([(20, "Genel", D(-3), D("1250.50"), None)], ["Satır okunamadı"] * 5))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--durum", str(y), "--cikti", c, "--tarih", "12.10.2026"]
            self.assertEqual(main.main(arg + ["--hedef", "80", "--duyarlilik", "0.75", "--azami", "10", "--rakip-bant", "20", "--pickup-gun", "3", "--projeksiyon", "7", "--yuvarla", "0"]), 0)
            for ek in (["--hedef", "0"], ["--duyarlilik", "-1"], ["--rakip-bant", "100"], ["--pickup-gun", "0"], ["--tarih", "x"], ["--yuvarla", "x"]):
                self.assertEqual(main.main(arg + ek), 2)
            self.assertEqual(main.main(["--durum", str(y), "--cikti", c, "--tarih", "01.11.2026"]), 1)
            self.assertEqual(main.main(["--durum", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Tarih;Kapasite\n1;2\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
