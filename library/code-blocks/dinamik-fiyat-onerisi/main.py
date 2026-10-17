"""
Dinamik Fiyat Önerisi — Gerisi Bende kod bloğu
Rezervasyon ve Gelir Yönetimi › Gelir Müdürü (Revenue Manager)

Doluluk, pick-up hızı ve rakip fiyatından tarih ve oda tipi bazında kural tabanlı, gerekçeli bir fiyat önerisi üretir.
  Günlük hız        = son pick-up ÷ --pickup-gun (7)                                              (pick-up penceresindeki günlük ortalama)
  Tahmini doluluk % = (satılan oda + günlük hız × min(kalan gün, --projeksiyon 14)) ÷ kapasite × 100
  Ayar %            = (tahmini doluluk − --hedef 85) × --duyarlilik 0,5;  ±--azami 20 ile sınırlanır
  Ham fiyat         = mevcut fiyat × (1 + ayar ÷ 100)
  Rakip bandı       = rakip fiyat × (1 ± --rakip-bant 15 ÷ 100):  artışta fiyat bandın üstüne, indirimde bandın altına çıkarılmaz; bant öneriyi ters yöne çevirmez
  Yuvarlama         = --yuvarla (10) katına;  banda takılan fiyat bandın içinde kalacak yöne yuvarlanır
  Taban / tavan     = satırda verilmişse son fiyat bu aralığa çekilir
Bu bir gelir yönetimi sistemi değildir: talep tahmini yapmaz, esnekliği ölçmez; şeffaf birkaç kuralı uygular. Öneriyi etkinlik, sezon, grup ve kanal bilginizle siz değerlendirirsiniz.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
from decimal import ROUND_CEILING, ROUND_FLOOR

SUTUN = {"tarih": ("tarih", "konaklama tarihi", "gun"), "oda": ("oda tipi", "tip", "oda tipi kodu"), "kapasite": ("kapasite", "satilabilir oda", "oda sayisi"), "satilan": ("satilan oda", "satilan", "rezerve oda", "otb"),
         "pickup": ("son 7 gun pick up", "pick up", "pickup", "son pick up", "son 7 gun pickup"), "mevcut": ("mevcut fiyat", "fiyat", "satis fiyati", "bar"), "rakip": ("rakip fiyat", "rakip ortalama fiyat", "rakip ortalama", "rakip"),
         "taban": ("taban fiyat", "taban", "asgari fiyat"), "tavan": ("tavan fiyat", "tavan", "azami fiyat")}
Z, YUZ = Decimal(0), Decimal(100)
AYAR = {"hedef": Decimal(85), "duyarlilik": Decimal("0.5"), "azami": Decimal(20), "rakip_bant": Decimal(15), "pickup_gun": 7, "projeksiyon": 14, "yuvarla": Decimal(10)}
TUR_RENK = {"Artır": YESIL, "İndir": SARI}


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("tarih", "kapasite", "satilan", "mevcut")):
        if not metin(r["tarih"]) and not metin(r["mevcut"]):
            continue
        t, o = tarih(r["tarih"]), metin(r["oda"]) or "Genel"
        d = {a: para(r[a]) if metin(r[a]) else None for a in ("kapasite", "satilan", "pickup", "mevcut", "rakip", "taban", "tavan")}
        bozuk = [a for a in d if d[a] is None and metin(r[a])] + [a for a in ("kapasite", "mevcut", "rakip", "taban", "tavan") if d[a] is not None and d[a] <= 0]
        if t is None or d["kapasite"] is None or d["satilan"] is None or d["mevcut"] is None or d["satilan"] < 0 or bozuk or (d["taban"] and d["tavan"] and d["taban"] > d["tavan"]):
            uy.append(uyari("Yüksek", "Satır okunamadı", f"{metin(r['tarih'])} · {o} (satır {r['_satir']})", f"Tarih, kapasite, satılan oda ve mevcut fiyat dolu ve geçerli olmalı; taban fiyat tavandan büyük olamaz"
                            + (f" (sorunlu alan: {', '.join(dict.fromkeys(bozuk))})" if bozuk else "") + ". Satır alınmadı."))
            continue
        l.append({"tarih": t, "oda": o} | d)
    return l, uy


def katina(x: Decimal, adim: Decimal, yon=ROUND_HALF_UP) -> Decimal:
    return (x / adim).quantize(Decimal(1), rounding=yon) * adim if adim > 0 else yuvarla(x)


def oner(satirlar: list[dict], bugun: date, **ayar) -> dict:
    a = AYAR | ayar
    if not satirlar:
        raise ValueError("Değerlendirilecek satır yok.")
    uy, liste = [], []
    for s in sorted(satirlar, key=lambda s: (s["tarih"], s["oda"])):
        kalan = (s["tarih"] - bugun).days
        if kalan < 0:
            uy.append(uyari("Bilgi", "Geçmiş tarih", f"{s['tarih']:%d.%m.%Y} · {s['oda']}", "Konaklama tarihi bugünden önce; öneri üretilmedi."))
            continue
        kap, mevcut = s["kapasite"], s["mevcut"]
        hiz = (s["pickup"] or Z) / a["pickup_gun"]
        tahmin = max((s["satilan"] + hiz * min(kalan, a["projeksiyon"])) / kap * 100, Z)
        ham = (tahmin - a["hedef"]) * a["duyarlilik"]
        ayar_y = max(min(ham, a["azami"]), -a["azami"])
        fiyat, gerekce, yon = mevcut * (1 + ayar_y / 100), [f"Tahmini doluluk %{yuvarla(tahmin)}, hedef %{a['hedef'].normalize():f}: ayar {yuvarla(ayar_y):+f}%" + (" (azami sınırda)" if ayar_y != ham else "")], ROUND_HALF_UP
        if s["rakip"]:
            alt, ust = s["rakip"] * (1 - a["rakip_bant"] / 100), s["rakip"] * (1 + a["rakip_bant"] / 100)
            if ayar_y > 0 and fiyat > max(mevcut, ust):
                fiyat, yon = max(mevcut, ust), ROUND_FLOOR
                gerekce.append(f"Rakip bandının üst sınırı {yuvarla(ust):,.2f}; artış sınırlandı")
            elif ayar_y < 0 and fiyat < min(mevcut, alt):
                fiyat, yon = min(mevcut, alt), ROUND_CEILING
                gerekce.append(f"Rakip bandının alt sınırı {yuvarla(alt):,.2f}; indirim sınırlandı")
        fiyat = katina(fiyat, a["yuvarla"], yon)
        if (ayar_y > 0 and fiyat < mevcut) or (ayar_y < 0 and fiyat > mevcut):
            fiyat = mevcut                                                                       # yuvarlama veya bant öneriyi ters yöne çeviremez
        if s["taban"] and fiyat < s["taban"]:
            fiyat = s["taban"]
            gerekce.append(f"Taban fiyat {s['taban']:,.2f} uygulandı")
        if s["tavan"] and fiyat > s["tavan"]:
            fiyat = s["tavan"]
            gerekce.append(f"Tavan fiyat {s['tavan']:,.2f} uygulandı")
        fiyat = yuvarla(fiyat)
        degisim = yuvarla((fiyat - mevcut) / mevcut * 100)
        tur = "Artır" if fiyat > mevcut else "İndir" if fiyat < mevcut else "Koru"
        if s["satilan"] >= kap:
            gerekce.append("Kapasite dolu; satışa açık oda yoksa öneri yalnız bilgi amaçlıdır")
        if s["satilan"] > kap:
            uy.append(uyari("Orta", "Satılan oda kapasiteden fazla", f"{s['tarih']:%d.%m.%Y} · {s['oda']}", f"Satılan {s['satilan']:f}, kapasite {kap:f}. Fazla rezervasyon bilinçli değilse kontrol edin."))
        if s["rakip"] and mevcut > s["rakip"] * (1 + a["rakip_bant"] / 100) and tur != "İndir":
            uy.append(uyari("Bilgi", "Mevcut fiyat rakip bandının üstünde", f"{s['tarih']:%d.%m.%Y} · {s['oda']}", f"Mevcut {mevcut:,.2f}, rakip {s['rakip']:,.2f}. Talep güçlü olduğu için indirim önerilmedi; konumlandırmanızı gözden geçirin."))
        liste.append(s | {"kalan": kalan, "doluluk": yuvarla(s["satilan"] / kap * 100), "hiz": yuvarla(hiz), "tahmin": yuvarla(tahmin), "ayar": yuvarla(ayar_y), "oneri": fiyat, "degisim": degisim, "tur": tur,
                          "rakip_fark": yuvarla((fiyat - s["rakip"]) / s["rakip"] * 100) if s["rakip"] else None, "gerekce": ". ".join(gerekce) + "."})
    if not liste:
        raise ValueError("Bütün satırlar geçmiş tarihli; öneri üretilecek gün yok.")
    say = {t: sum(1 for x in liste if x["tur"] == t) for t in ("Artır", "İndir", "Koru")}
    return {"oneriler": liste, "say": say, "ayar": a, "bugun": bugun, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb, a = Workbook(), r["ayar"]
    ws = wb.active
    ws.title = "Fiyat Önerisi"
    _baslik(ws, ["Tarih", "Oda Tipi", "Kalan Gün", "Kapasite", "Satılan", "Doluluk %", "Günlük Pick-up", "Tahmini Doluluk %", "Mevcut Fiyat", "Rakip Fiyat", "Ayar %", "Önerilen Fiyat", "Değişim %", "Öneri", "Rakibe Göre %", "Gerekçe",
                 "Karar (siz doldurun)"], (12, 16, 8, 9, 9, 9, 9, 10, 13, 13, 8, 13, 9, 9, 10, 70, 22), "C2")
    for x in r["oneriler"]:
        _satir(ws, [x["tarih"], x["oda"], x["kalan"], x["kapasite"], x["satilan"], x["doluluk"], x["hiz"], x["tahmin"], x["mevcut"], x["rakip"], x["ayar"], x["oneri"], x["degisim"], x["tur"], x["rakip_fark"], x["gerekce"], ""],
               para_s=(6, 7, 8, 9, 10, 11, 12, 13, 15), tarih_s=(1,), dolgu={14: TUR_RENK.get(x["tur"])}, sar=(16,))
    ws = wb.create_sheet("Kurallar")
    _baslik(ws, ["Kural", "Değer", "Açıklama"], (26, 12, 100))
    for ad, deger, ac in (("Bugün", f"{r['bugun']:%d.%m.%Y}", "Kalan gün bu tarihe göre hesaplanır"), ("Hedef doluluk %", a["hedef"], "Tahmini doluluk bunun üstündeyse fiyat artırılır, altındaysa indirilir"),
                          ("Duyarlılık", a["duyarlilik"], "Her 1 puan doluluk farkı için fiyat ayarı (%)"), ("Azami ayar %", a["azami"], "Tek seferde önerilecek en büyük artış veya indirim"),
                          ("Rakip bandı %", a["rakip_bant"], "Öneri rakip fiyatın bu kadar üstünü veya altını aşacak şekilde değişmez"), ("Pick-up penceresi (gün)", a["pickup_gun"], "Dosyadaki pick-up değerinin kapsadığı gün sayısı"),
                          ("Projeksiyon (gün)", a["projeksiyon"], "Günlük hızın en çok kaç gün süreceği varsayılır"), ("Yuvarlama", a["yuvarla"], "Önerilen fiyat bu değerin katına yuvarlanır")):
        _satir(ws, [ad, deger, ac], sar=(3,))
    ws.append([])
    ws.append(["Kural tabanlı bir başlangıç önerisidir; talep tahmini yapmaz. Etkinlik, sezon, grup, kanal ve sözleşmeli fiyat bilginizle birlikte değerlendirin."])
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, **ayar) -> dict:
    satirlar, uy = oku(yol)
    r = oner(satirlar, bugun, **ayar)
    uy += r["uyarilar"]
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Doluluk, pick-up hızı ve rakip fiyatından kural tabanlı fiyat önerisi üretir.")
    p.add_argument("--durum", type=Path, default=ORNEK / "durum.csv", help="Tarih, Oda Tipi, Kapasite, Satılan Oda, Son 7 Gün Pick-up, Mevcut Fiyat, Rakip Fiyat, Taban Fiyat, Tavan Fiyat")
    p.add_argument("--tarih", help="Bugünün tarihi (GG.AA.YYYY); verilmezse sistem tarihi")
    p.add_argument("--hedef", default="85", help="Hedef doluluk %% (varsayılan 85)")
    p.add_argument("--duyarlilik", default="0,5", help="Her 1 puan doluluk farkı için fiyat ayarı %% (varsayılan 0,5)")
    p.add_argument("--azami", default="20", help="Azami artış veya indirim %% (varsayılan 20)")
    p.add_argument("--rakip-bant", default="15", help="Rakip fiyat bandı %% (varsayılan 15)")
    p.add_argument("--pickup-gun", type=int, default=7, help="Pick-up değerinin kapsadığı gün sayısı (varsayılan 7)")
    p.add_argument("--projeksiyon", type=int, default=14, help="Günlük hızın süreceği varsayılan en çok gün (varsayılan 14)")
    p.add_argument("--yuvarla", default="10", help="Fiyat bu değerin katına yuvarlanır; 0 ise kuruşa (varsayılan 10)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "fiyat_onerisi.xlsx")
    a = p.parse_args(argv)
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if a.durum == ORNEK / "durum.csv" else date.today()
    h, d, az, rb, yv = (para(x) for x in (a.hedef, a.duyarlilik, a.azami, a.rakip_bant, a.yuvarla))
    if bugun is None or None in (h, d, az, rb, yv) or not 0 < h <= 100 or d < 0 or az < 0 or not 0 <= rb < 100 or yv < 0 or a.pickup_gun <= 0 or a.projeksiyon < 0:
        print("[X] --tarih GG.AA.YYYY; --hedef 0–100; --duyarlilik, --azami, --yuvarla sıfır veya pozitif; --rakip-bant 0–100; --pickup-gun pozitif; --projeksiyon sıfır veya pozitif olmalı")
        return 2
    if not a.durum.exists():
        print(f"[X] Dosya bulunamadı: {a.durum}")
        return 1
    try:
        r = calistir(a.durum, a.cikti, bugun, hedef=h, duyarlilik=d, azami=az, rakip_bant=rb, pickup_gun=a.pickup_gun, projeksiyon=a.projeksiyon, yuvarla=yv)
    except ValueError as hata:
        print(f"[X] {hata}")
        return 1
    s = r["say"]
    print(f"[OK] {len(r['oneriler'])} satır · {s['Artır']} artır · {s['İndir']} indir · {s['Koru']} koru")
    for x in r["oneriler"]:
        if x["tur"] != "Koru":
            print(f"[!] {x['tarih']:%d.%m.%Y} {x['oda']}: {x['mevcut']:,.2f} → {x['oneri']:,.2f} ({x['degisim']:+f}%), tahmini doluluk %{x['tahmin']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
