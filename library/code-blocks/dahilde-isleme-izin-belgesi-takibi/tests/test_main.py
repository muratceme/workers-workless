import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)
IHR, ITH = "İhracat", "İthalat"


def d(x):
    return None if x is None else D(str(x))


def s(belge, tur=IHR, madde="M", miktar=100, tutar=1000, sure_sonu=date(2027, 6, 1)):
    return {"belge": belge, "belge_tarihi": date(2025, 1, 1), "sure_sonu": sure_sonu, "tur": tur, "madde": madde, "birim": "kg", "miktar": d(miktar), "tutar": d(tutar), "satir": 2}


def h(belge, tur=IHR, madde="M", miktar=40, tutar=400, t=date(2026, 1, 5)):
    return {"belge": belge, "tarih": t, "tur": tur, "madde": madde, "beyanname": "BY", "miktar": d(miktar), "tutar": d(tutar), "satir": 2}


class DegerlendirTesti(unittest.TestCase):
    def test_durum_sinirlari(self):
        g = lambda n: BUGUN + timedelta(days=n)   # noqa: E731
        sonlar = {"bugün": g(0), "dün": g(-1), "60": g(60), "61": g(61), "kapatma_bugün": date(2026, 7, 12), "kapatma_dün": date(2026, 7, 11), "yok": None}
        tamam = {"t_uzak": g(200), "t_61": date(2026, 9, 12), "t_60": date(2026, 9, 11), "t_gecmis": date(2026, 7, 11)}
        r = main.degerlendir([s(k, sure_sonu=v) for k, v in (sonlar | tamam).items()], [h(k) for k in sonlar] + [h(k, miktar=100, tutar=1000) for k in tamam], BUGUN)
        o = {x["no"]: x for x in r["belgeler"]}
        self.assertEqual({k: o[k]["durum"] for k in o},
                         {"bugün": "Süre yaklaşıyor", "dün": "Süresi doldu, taahhüt açık", "60": "Süre yaklaşıyor", "61": "Devam ediyor", "kapatma_bugün": "Süresi doldu, taahhüt açık", "kapatma_dün": "Kapatma süresi geçmiş",
                          "yok": "Süre sonu tanımsız", "t_uzak": "Taahhüt tamamlandı, kapatma bekliyor", "t_61": "Taahhüt tamamlandı, kapatma bekliyor", "t_60": "Taahhüt tamamlandı, kapatma bekliyor", "t_gecmis": "Kapatma süresi geçmiş"})
        self.assertEqual([(o[k]["kalan_gun"], o[k]["kapatma_son"], o[k]["kapatma_kalan"]) for k in ("dün", "kapatma_bugün", "t_60", "yok")],
                         [(-1, date(2027, 1, 11), 91), (-92, date(2026, 10, 12), 0), (-31, date(2026, 12, 11), 60), (None, None, None)])
        self.assertEqual([x["no"] for x in r["belgeler"]], ["kapatma_dün", "t_gecmis", "kapatma_bugün", "dün", "bugün", "60", "t_60", "t_61", "t_uzak", "61", "yok"])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         sorted([("Orta", "Süre yaklaşıyor", "bugün"), ("Yüksek", "Süresi doldu, taahhüt açık", "dün"), ("Orta", "Süre yaklaşıyor", "60"), ("Yüksek", "Süresi doldu, taahhüt açık", "kapatma_bugün"),
                                 ("Yüksek", "Kapatma başvuru süresi geçmiş", "kapatma_dün"), ("Yüksek", "Süre sonu yok", "yok"), ("Orta", "Kapatma başvurusu yaklaşıyor", "t_60"), ("Yüksek", "Kapatma başvuru süresi geçmiş", "t_gecmis")]))
        a = {u["kim"]: u["aciklama"] for u in r["uyarilar"]}
        self.assertIn("Açık taahhüt: M: 60 kg.", a["dün"])
        self.assertIn("Açık taahhüt: M: 60 kg.", a["kapatma_dün"])
        self.assertNotIn("Açık taahhüt", a["t_gecmis"])
        f = lambda **kw: main.degerlendir([s("a", sure_sonu=date(2026, 12, 31))], [h("a")], BUGUN, **kw)["belgeler"][0]   # noqa: E731
        self.assertEqual((f()["durum"], f(gun=80)["durum"], f(kapatma_ay=1)["kapatma_son"], f(kapatma_ay=2)["kapatma_son"]), ("Devam ediyor", "Süre yaklaşıyor", date(2027, 1, 31), date(2027, 2, 28)))
        with self.assertRaises(ValueError):
            main.degerlendir([], [], BUGUN)

    def test_gerceklesme(self):
        satirlar = [s("B", ITH, "Ham", 100, 500), s("B", IHR, "Mamul", 80, 2000), s("B", IHR, "Mamul 2", None, 1000), s("B", IHR, "mamul", 20, 500), s("C", ITH, "Ham", None, 300, None), s("C", IHR, "X", 0, 100, None)]
        hareketler = [h("B", ITH, "ham", "100.5", 510), h("b", IHR, "MAMUL", 60, 1500), h("B", IHR, "Mamul", 40, 1100, None), h("B", IHR, "Mamul 2", None, "999.99"), h("B", IHR, "Yok", 5, 50), h("B", IHR, "Mamul", 1, 1, date(2027, 6, 2)),
                      h("Z", IHR, "Mamul", 1, 7), h("C", ITH, "Ham", None, 300), h("C", IHR, "X", 3, 100)]
        r = main.degerlendir(satirlar, hareketler, BUGUN)
        self.assertEqual([(e["belge"], e["tur"], e["madde"], e["miktar"], e["gercek_miktar"], e["kalan_miktar"], e["tutar"], e["gercek_tutar"], e["kalan_tutar"], e["oran"], e["tamam"]) for e in r["satirlar"]],
                         [("B", ITH, "Ham", D(100), D("100.5"), D("-0.5"), D(500), D(510), D(-10), D("100.50"), True), ("B", IHR, "Mamul", D(100), D(100), D(0), D(2500), D(2600), D(-100), D("100.00"), True),
                          ("B", IHR, "Mamul 2", None, D(0), None, D(1000), D("999.99"), D("0.01"), D("100.00"), False), ("C", ITH, "Ham", None, D(0), None, D(300), D(300), D(0), D("100.00"), True),
                          ("C", IHR, "X", D(0), D(3), None, D(100), D(100), D(0), D("100.00"), True)])
        b, c = r["belgeler"]
        self.assertEqual((b["no"], b["ihr_taahhut"], b["ihr_gercek"], b["ihr_kalan"], b["ihr_oran"], b["ith_izin"], b["ith_gercek"], b["ith_kalan"], b["ith_oran"], b["dko"], b["tamam"], b["durum"]),
                         ("B", D(3500), D("3649.99"), D("-149.99"), D("104.29"), D(500), D(510), D(-10), D("102.00"), D("13.97"), False, "Devam ediyor"))
        self.assertEqual((c["no"], c["tamam"], c["durum"], c["dko"]), ("C", True, "Süre sonu tanımsız", D("300.00")))
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         sorted([("Orta", "Belge satırı tekrarlanıyor", "B · mamul"), ("Yüksek", "İthalat izni aşıldı", "B · Ham"), ("Orta", "Belgede olmayan madde", "B · BY"), ("Orta", "Süre sonundan sonra hareket", "B · BY"),
                                 ("Orta", "Tanımsız belge", "Z · BY"), ("Yüksek", "Süre sonu yok", "C")]))
        r = main.degerlendir([s("A", ITH, miktar=100), s("K", ITH, miktar=None, tutar=100)], [h("A", ITH, miktar=100), h("K", ITH, miktar=None, tutar="100.01")], BUGUN)
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]), [("İhracat taahhüdü satırı yok", "A"), ("İhracat taahhüdü satırı yok", "K"), ("İthalat izni aşıldı", "K · M")])
        self.assertEqual([(x["tamam"], x["ihr_oran"], x["dko"]) for x in r["belgeler"]], [(False, None, None)] * 2)

    def test_ay_ekle(self):
        self.assertEqual([main.ay_ekle(date(*t), n) for t, n in (((2026, 11, 30), 3), ((2027, 11, 30), 3), ((2026, 10, 31), 1), ((2026, 12, 15), 3), ((2026, 5, 5), 0))],
                         [date(2027, 2, 28), date(2028, 2, 29), date(2026, 11, 30), date(2027, 3, 15), date(2026, 5, 5)])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "belge_satirlari.csv", ORNEK / "hareketler.csv", cikti, BUGUN)
            self.assertEqual([(o["no"], o["durum"], o["kalan_gun"], o["kapatma_son"], o["ihr_taahhut"], o["ihr_gercek"], o["ihr_oran"], o["ihr_kalan"], o["ith_izin"], o["ith_gercek"], o["ith_oran"], o["ith_kalan"], o["dko"])
                              for o in r["belgeler"]],
                             [("D1-2025-0177", "Süresi doldu, taahhüt açık", -72, date(2026, 11, 1), D(81000), D(61000), D("75.31"), D(20000), D(42000), D(42000), D("100.00"), D(0), D("68.85")),
                              ("D1-2025-0101", "Süre yaklaşıyor", 34, date(2027, 2, 15), D(315000), D(191000), D("60.63"), D(124000), D(150000), D(136000), D("90.67"), D(14000), D("71.20")),
                              ("D1-2025-0050", "Taahhüt tamamlandı, kapatma bekliyor", -41, date(2026, 12, 1), D(60000), D(60500), D("100.83"), D(-500), D(16000), D(16400), D("102.50"), D(-400), D("27.11")),
                              ("D1-2026-0042", "Devam ediyor", 121, date(2027, 5, 10), D(288000), D(82000), D("28.47"), D(206000), D(200000), D(100000), D("50.00"), D(100000), D("121.95"))])
            self.assertEqual([(e["madde"], e["gercek_miktar"], e["kalan_miktar"], e["oran"], e["tamam"]) for e in r["satirlar"]],
                             [("Polietilen Granül", D(30000), D(0), D("100.00"), True), ("Plastik Ambalaj", D(20000), D(7000), D("74.07"), False), ("Pamuk İpliği", D(45000), D(5000), D("90.00"), False),
                              ("Örme Kumaş", D(27000), D(18000), D("60.00"), False), ("Çelik Sac", D(20500), D(-500), D("102.50"), True), ("Metal Raf", D(4000), D(0), D("100.00"), True),
                              ("Alüminyum Külçe", D(40000), D(40000), D("50.00"), False), ("Alüminyum Profil", D(20000), D(52000), D("27.78"), False)])
            self.assertEqual(r["say"], {"Kapatma süresi geçmiş": 0, "Süresi doldu, taahhüt açık": 1, "Süre yaklaşıyor": 1, "Taahhüt tamamlandı, kapatma bekliyor": 1, "Devam ediyor": 1, "Süre sonu tanımsız": 0})
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Yüksek", "Süresi doldu, taahhüt açık", "D1-2025-0177"), ("Yüksek", "İthalat izni aşıldı", "D1-2025-0050 · Çelik Sac"), ("Orta", "Kapatma başvurusu yaklaşıyor", "D1-2025-0050"),
                              ("Orta", "Süre sonundan sonra hareket", "D1-2025-0177 · EX-0004"), ("Orta", "Süre yaklaşıyor", "D1-2025-0101"), ("Orta", "Tanımsız belge", "D1-2024-0999 · EX-0008")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Belgeler", "Satırlar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            by, hy = Path(tmp) / "b.csv", Path(tmp) / "h.csv"
            by.write_text("DİİB No;Başlangıç Tarihi;Bitiş Tarihi;İşlem;Eşya;Birim;Miktar;Değer USD\nD1;01.01.2026;31.12.2026;İTHAL;İplik;kg;1.250,5;12.500\nD1;;;ihracat;Kumaş;;;7.000,25\nD1;;;Transit;X;;1;1\nD1;;;İhracat;Y;;;\n"
                          "D1;;;İhracat;Z;;-1;5\n;;;İhracat;Q;;1;1\n", encoding="utf-8")
            hy.write_text("Belge No;Tescil Tarihi;Tür;Ürün;Tescil No;Miktar;Tutar\nD1;05.03.2026;İthalat;İplik;IM1;500,25;5000\nD1;x;Diğer;İplik;IM2;1;1\n", encoding="utf-8")
            l, uy = main.oku_belgeler(by)
            self.assertEqual(([(x["belge"], x["belge_tarihi"], x["sure_sonu"], x["tur"], x["madde"], x["birim"], x["miktar"], x["tutar"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("D1", date(2026, 1, 1), date(2026, 12, 31), ITH, "İplik", "kg", D("1250.5"), D(12500)), ("D1", None, None, IHR, "Kumaş", "", None, D("7000.25"))], [("Belge satırı okunamadı", "D1")] * 3))
            l, uy = main.oku_hareketler(hy)
            self.assertEqual(([(x["belge"], x["tarih"], x["tur"], x["madde"], x["beyanname"], x["miktar"], x["tutar"]) for x in l], [u["tur"] for u in uy]),
                             ([("D1", date(2026, 3, 5), ITH, "İplik", "IM1", D("500.25"), D(5000))], ["Hareket okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--belgeler", str(by), "--hareketler", str(hy), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "12.10.2026", "--gun", "90", "--kapatma-ay", "1"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + x) for x in (["--tarih", "x"], ["--gun", "-1"], ["--kapatma-ay", "25"])] + [main.main(["--belgeler", str(by)])], [2] * 4)
            self.assertEqual(main.main(["--belgeler", str(by), "--hareketler", str(Path(tmp) / "yok.csv")]), 1)
            hy.write_text("Belge No;Tarih\nA;1\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            hy.write_text("Belge No;Tür;Tutar\nD1;İhracat;5\n", encoding="utf-8")
            by.write_text("Belge No;Tür;Tutar\n;İhracat;5\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
