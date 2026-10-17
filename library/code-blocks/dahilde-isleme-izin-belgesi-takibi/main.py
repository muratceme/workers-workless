"""
Dahilde İşleme İzin Belgesi Takibi — Gerisi Bende kod bloğu
Dış Ticaret › Dış Ticaret Uzmanı

DİİB kapsamındaki ithalat izinlerini ve ihracat taahhütlerini, kalan miktarları ve süre sonunu izler.
  Gerçekleşen               = belge süresi içindeki hareketlerin toplamı (madde ve tür bazında)
  Kalan ihracat taahhüdü    = taahhüt − gerçekleşen ihracat;   kalan ithalat hakkı = izin − gerçekleşen ithalat
  Gerçekleşme %             = gerçekleşen ÷ taahhüt (ya da izin) × 100   (satırda miktar varsa miktar, yoksa tutar üzerinden)
  Taahhüt tamam             = belgenin bütün ihracat satırlarında gerçekleşen ≥ taahhüt
  Döviz kullanım oranı %    = gerçekleşen ithalat tutarı ÷ gerçekleşen ihracat tutarı × 100   (bilgi amaçlı)
  Kapatma başvurusu son gün = belge süre sonu + 3 ay   (--kapatma-ay ile değişir)
  Durum                     = Kapatma süresi geçmiş · Süresi doldu, taahhüt açık · Süre yaklaşıyor (60 gün) ·
                              Taahhüt tamamlandı, kapatma bekliyor · Devam ediyor
Süre sonundan sonraki hareketler gerçekleşmeye katılmaz ve uyarılır; ek süre aldıysanız süre sonunu güncelleyin.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
import calendar

B_SUTUN = {"belge": ("belge no", "diib no", "belge numarasi"), "belge_tarihi": ("belge tarihi", "baslangic tarihi"), "sure_sonu": ("sure sonu", "belge sure sonu", "bitis tarihi", "sure sonu tarihi"), "tur": ("tur", "islem"),
           "madde": ("madde", "madde adi", "esya", "urun", "gtip"), "birim": ("birim",), "miktar": ("miktar",), "tutar": ("tutar usd", "tutar", "deger usd", "deger")}
H_SUTUN = {"belge": ("belge no", "diib no", "belge numarasi"), "tarih": ("tarih", "beyanname tarihi", "tescil tarihi"), "tur": ("tur", "islem"), "madde": ("madde", "madde adi", "esya", "urun", "gtip"),
           "beyanname": ("beyanname no", "beyanname", "tescil no"), "miktar": ("miktar",), "tutar": ("tutar usd", "tutar", "deger usd", "deger")}
TUR = {"ithalat": "İthalat", "ithal": "İthalat", "ihracat": "İhracat", "ihrac": "İhracat"}
SIRA = {"Kapatma süresi geçmiş": 0, "Süresi doldu, taahhüt açık": 1, "Süre yaklaşıyor": 2, "Taahhüt tamamlandı, kapatma bekliyor": 3, "Devam ediyor": 4, "Süre sonu tanımsız": 5}
RENK_D = {"Kapatma süresi geçmiş": KIRMIZI, "Süresi doldu, taahhüt açık": KIRMIZI, "Süre yaklaşıyor": SARI, "Taahhüt tamamlandı, kapatma bekliyor": YESIL, "Devam ediyor": None, "Süre sonu tanımsız": MAVI}
Z, YUZ = Decimal(0), Decimal(100)


def ay_ekle(t: date, n: int) -> date:
    y, a = divmod(t.year * 12 + t.month - 1 + n, 12)
    return date(y, a + 1, min(t.day, calendar.monthrange(y, a + 1)[1]))


def olcu(x) -> Decimal | None:
    """Miktar okur: virgül ya da tek nokta ondalık ayırıcıdır; '1.250,5' binlik ayırıcılıdır."""
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x or ""))
    try:
        return Decimal(s.replace(".", "").replace(",", ".") if "," in s else s) if s else None
    except InvalidOperation:
        return None


def _oku(yol: Path, sozluk: dict, tarih_alanlari: tuple, ne: str) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, sozluk, ("belge", "tur")):
        belge = metin(r["belge"])
        if not belge:
            continue
        tur, miktar, tutar = TUR.get(katla(r["tur"])), olcu(r["miktar"]), para(r["tutar"])
        if tur is None or (miktar is None and tutar is None) or (miktar is not None and miktar < 0) or (tutar is not None and tutar < 0):
            uy.append(uyari("Yüksek", f"{ne} okunamadı", belge, f"{yol.name} satır {r['_satir']}: tür '{metin(r['tur'])}' (İthalat / İhracat olmalı), miktar '{metin(r['miktar'])}', tutar '{metin(r['tutar'])}'. Satır alınmadı."))
            continue
        l.append({a: metin(r[a]) for a in sozluk if a not in tarih_alanlari + ("tur", "miktar", "tutar")} | {a: tarih(r[a]) for a in tarih_alanlari} | {"tur": tur, "miktar": miktar, "tutar": tutar, "satir": r["_satir"]})
    return l, uy


def oku_belgeler(yol: Path) -> tuple[list[dict], list[dict]]:
    return _oku(yol, B_SUTUN, ("belge_tarihi", "sure_sonu"), "Belge satırı")


def oku_hareketler(yol: Path) -> tuple[list[dict], list[dict]]:
    return _oku(yol, H_SUTUN, ("tarih",), "Hareket")


def _oran(pay: Decimal, payda: Decimal | None) -> Decimal | None:
    return yuvarla(pay * YUZ / payda) if payda else None


def degerlendir(satirlar: list[dict], hareketler: list[dict], bugun: date, gun: int = 60, kapatma_ay: int = 3) -> dict:
    if not satirlar:
        raise ValueError("Belge satırı yok.")
    uy, belgeler = [], {}
    for s in satirlar:
        b = belgeler.setdefault(katla(s["belge"]), {"no": s["belge"], "belge_tarihi": None, "sure_sonu": None, "satirlar": {}, "disi": {"İthalat": Z, "İhracat": Z}})
        b["belge_tarihi"], b["sure_sonu"] = b["belge_tarihi"] or s["belge_tarihi"], b["sure_sonu"] or s["sure_sonu"]
        k = (s["tur"], katla(s["madde"]))
        if k in b["satirlar"]:
            uy.append(uyari("Orta", "Belge satırı tekrarlanıyor", f"{s['belge']} · {s['madde']}", f"Aynı belge, tür ve madde birden çok satırda (satır {s['satir']}); miktar ve tutarlar toplandı."))
            e = b["satirlar"][k]
            e["miktar"], e["tutar"] = (e["miktar"] or Z) + (s["miktar"] or Z) if e["miktar"] is not None or s["miktar"] is not None else None, (e["tutar"] or Z) + (s["tutar"] or Z)
        else:
            b["satirlar"][k] = {"belge": s["belge"], "tur": s["tur"], "madde": s["madde"], "birim": s["birim"], "miktar": s["miktar"], "tutar": s["tutar"] or Z, "gercek_miktar": Z, "gercek_tutar": Z}
    for h in hareketler:
        b = belgeler.get(katla(h["belge"]))
        kim = f"{h['belge']} · {h['beyanname'] or h['madde']}".strip(" ·")
        if b is None:
            uy.append(uyari("Orta", "Tanımsız belge", kim, f"Hareket dosyasında geçen belge, belge dosyasında yok ({h['tur']}, {h['tutar'] or Z:,.2f}). Hareket hesaba katılmadı."))
            continue
        if b["sure_sonu"] and h["tarih"] and h["tarih"] > b["sure_sonu"]:
            uy.append(uyari("Orta", "Süre sonundan sonra hareket", kim, f"{h['tur']} tarihi {h['tarih']:%d.%m.%Y}, belge süre sonu {b['sure_sonu']:%d.%m.%Y}. Gerçekleşmeye katılmadı; ek süre alındıysa süre sonunu güncelleyin."))
            continue
        e = b["satirlar"].get((h["tur"], katla(h["madde"])))
        if e is None:
            b["disi"][h["tur"]] += h["tutar"] or Z
            uy.append(uyari("Orta", "Belgede olmayan madde", kim, f"{h['tur']} hareketindeki '{h['madde'] or 'madde yazılmamış'}' belge satırlarında yok. Tutar belge toplamına eklendi ama hiçbir satıra işlenmedi."))
            continue
        e["gercek_miktar"] += h["miktar"] or Z
        e["gercek_tutar"] += h["tutar"] or Z
    sonuc, tum_satirlar = [], []
    for b in belgeler.values():
        for e in b["satirlar"].values():
            miktarli = e["miktar"] is not None and e["miktar"] > 0
            e["kalan_miktar"] = e["miktar"] - e["gercek_miktar"] if miktarli else None
            e["kalan_tutar"] = e["tutar"] - e["gercek_tutar"]
            e["oran"] = _oran(e["gercek_miktar"], e["miktar"]) if miktarli else _oran(e["gercek_tutar"], e["tutar"])
            e["tamam"] = (e["gercek_miktar"] >= e["miktar"]) if miktarli else (e["tutar"] > 0 and e["gercek_tutar"] >= e["tutar"])
            asim = e["gercek_miktar"] > e["miktar"] if miktarli else e["gercek_tutar"] > e["tutar"] > 0
            if e["tur"] == "İthalat" and asim:
                uy.append(uyari("Yüksek", "İthalat izni aşıldı", f"{e['belge']} · {e['madde']}", f"İzin {e['miktar']:f} {e['birim']}, gerçekleşen {e['gercek_miktar']:f} {e['birim']}." if miktarli else
                                f"İzin {e['tutar']:,.2f}, gerçekleşen {e['gercek_tutar']:,.2f}."))
            tum_satirlar.append(e)
        ihr, ith = [e for e in b["satirlar"].values() if e["tur"] == "İhracat"], [e for e in b["satirlar"].values() if e["tur"] == "İthalat"]
        o = {"no": b["no"], "belge_tarihi": b["belge_tarihi"], "sure_sonu": b["sure_sonu"], "ihr_taahhut": sum((e["tutar"] for e in ihr), Z), "ihr_gercek": sum((e["gercek_tutar"] for e in ihr), Z) + b["disi"]["İhracat"],
             "ith_izin": sum((e["tutar"] for e in ith), Z), "ith_gercek": sum((e["gercek_tutar"] for e in ith), Z) + b["disi"]["İthalat"]}
        o["ihr_kalan"], o["ith_kalan"] = o["ihr_taahhut"] - o["ihr_gercek"], o["ith_izin"] - o["ith_gercek"]
        o["ihr_oran"], o["ith_oran"], o["dko"] = _oran(o["ihr_gercek"], o["ihr_taahhut"]), _oran(o["ith_gercek"], o["ith_izin"]), _oran(o["ith_gercek"], o["ihr_gercek"])
        o["tamam"] = bool(ihr) and all(e["tamam"] for e in ihr)
        o["kalan_gun"] = (o["sure_sonu"] - bugun).days if o["sure_sonu"] else None
        o["kapatma_son"] = ay_ekle(o["sure_sonu"], kapatma_ay) if o["sure_sonu"] else None
        o["kapatma_kalan"] = (o["kapatma_son"] - bugun).days if o["kapatma_son"] else None
        eksik = "; ".join(f"{e['madde']}: {e['kalan_miktar']:f} {e['birim']}".strip() if e["kalan_miktar"] is not None else f"{e['madde']}: {e['kalan_tutar']:,.2f}" for e in ihr if not e["tamam"])
        if not ihr:
            uy.append(uyari("Orta", "İhracat taahhüdü satırı yok", o["no"], "Belgede ihracat satırı tanımlı değil; taahhüt izlenemiyor."))
        if o["sure_sonu"] is None:
            o["durum"] = "Süre sonu tanımsız"
            uy.append(uyari("Yüksek", "Süre sonu yok", o["no"], "Belgenin süre sonu yazılmamış; süre ve kapatma takibi yapılamıyor."))
        elif o["kapatma_kalan"] < 0:
            o["durum"] = "Kapatma süresi geçmiş"
            uy.append(uyari("Yüksek", "Kapatma başvuru süresi geçmiş", o["no"], f"Süre sonu {o['sure_sonu']:%d.%m.%Y}, kapatma başvurusu için son gün {o['kapatma_son']:%d.%m.%Y} ({-o['kapatma_kalan']} gün geçti). "
                            + (f"Açık taahhüt: {eksik}. " if eksik else "") + "Belge kapatıldıysa listeden çıkarın."))
        elif o["tamam"]:
            o["durum"] = "Taahhüt tamamlandı, kapatma bekliyor"
            if o["kapatma_kalan"] <= gun:
                uy.append(uyari("Orta", "Kapatma başvurusu yaklaşıyor", o["no"], f"Taahhüt tamamlandı. Kapatma başvurusu için son gün {o['kapatma_son']:%d.%m.%Y} ({o['kapatma_kalan']} gün kaldı)."))
        elif o["kalan_gun"] < 0:
            o["durum"] = "Süresi doldu, taahhüt açık"
            uy.append(uyari("Yüksek", "Süresi doldu, taahhüt açık", o["no"], f"Süre {o['sure_sonu']:%d.%m.%Y} tarihinde doldu ({-o['kalan_gun']} gün önce). Açık taahhüt: {eksik or 'ihracat satırı yok'}. "
                            f"Kapatma başvurusu için son gün {o['kapatma_son']:%d.%m.%Y}."))
        elif o["kalan_gun"] <= gun:
            o["durum"] = "Süre yaklaşıyor"
            uy.append(uyari("Orta", "Süre yaklaşıyor", o["no"], f"Süre sonu {o['sure_sonu']:%d.%m.%Y} ({o['kalan_gun']} gün kaldı). Açık taahhüt: {eksik or 'ihracat satırı yok'}."))
        else:
            o["durum"] = "Devam ediyor"
        sonuc.append(o)
    sonuc.sort(key=lambda o: (SIRA[o["durum"]], o["sure_sonu"] or date.max, katla(o["no"])))
    sira = {katla(o["no"]): i for i, o in enumerate(sonuc)}
    tum_satirlar.sort(key=lambda e: (sira[katla(e["belge"])], e["tur"] != "İthalat", katla(e["madde"])))
    return {"belgeler": sonuc, "satirlar": tum_satirlar, "say": {d: sum(1 for o in sonuc if o["durum"] == d) for d in SIRA}, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, bugun: date) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Belgeler"
    _baslik(ws, ["Durum", "Belge No", "Belge Tarihi", "Süre Sonu", "Kalan Gün", "Kapatma Son Gün", "İhracat Taahhüdü", "Gerçekleşen İhracat", "İhracat %", "Kalan Taahhüt", "İthalat İzni", "Gerçekleşen İthalat", "İthalat %",
                 "Kalan İthalat Hakkı", "Döviz Kullanım Oranı %"], (34, 18, 12, 12, 9, 13, 16, 16, 10, 16, 16, 16, 10, 16, 12))
    for o in r["belgeler"]:
        _satir(ws, [o["durum"], o["no"], o["belge_tarihi"], o["sure_sonu"], o["kalan_gun"], o["kapatma_son"], o["ihr_taahhut"], o["ihr_gercek"], o["ihr_oran"], o["ihr_kalan"], o["ith_izin"], o["ith_gercek"], o["ith_oran"], o["ith_kalan"],
                    o["dko"]], para_s=(7, 8, 9, 10, 11, 12, 13, 14, 15), tarih_s=(3, 4, 6), dolgu={1: RENK_D[o["durum"]], 14: KIRMIZI if o["ith_kalan"] < 0 else None})
    ws.append([])
    ws.append([f"Durum tarihi: {bugun:%d.%m.%Y}. Tutarlar belgedeki para birimindedir. Döviz kullanım oranı bilgi amaçlıdır; belgenizde öngörülen oranla ayrıca karşılaştırın."])
    ws = wb.create_sheet("Satırlar")
    _baslik(ws, ["Belge No", "Tür", "Madde", "Birim", "Belge Miktarı", "Gerçekleşen Miktar", "Kalan Miktar", "Belge Tutarı", "Gerçekleşen Tutar", "Kalan Tutar", "Gerçekleşme %"], (18, 10, 30, 8, 14, 14, 14, 16, 16, 16, 12))
    for e in r["satirlar"]:
        _satir(ws, [e["belge"], e["tur"], e["madde"], e["birim"], e["miktar"], e["gercek_miktar"] if e["miktar"] is not None else None, e["kalan_miktar"], e["tutar"], e["gercek_tutar"], e["kalan_tutar"], e["oran"]],
               para_s=(5, 6, 7, 8, 9, 10, 11), dolgu={11: YESIL if e["tamam"] and e["tur"] == "İhracat" else None, 7: KIRMIZI if e["tur"] == "İthalat" and e["kalan_miktar"] is not None and e["kalan_miktar"] < 0 else None})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(b_yol: Path, h_yol: Path, cikti: Path, bugun: date, gun: int = 60, kapatma_ay: int = 3) -> dict:
    satirlar, uy1 = oku_belgeler(b_yol)
    hareketler, uy2 = oku_hareketler(h_yol)
    r = degerlendir(satirlar, hareketler, bugun, gun, kapatma_ay)
    r["uyarilar"] = sirala(uy1 + uy2 + r["uyarilar"])
    rapor_yaz(cikti, r, bugun)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="DİİB kapsamındaki ithalat izinlerini, ihracat taahhütlerini, kalan miktarları ve süre sonunu izler.")
    p.add_argument("--belgeler", type=Path, help="Belge satırları: Belge No, Belge Tarihi, Süre Sonu, Tür (İthalat / İhracat), Madde, Birim, Miktar, Tutar")
    p.add_argument("--hareketler", type=Path, help="Beyannameler: Belge No, Tarih, Tür, Madde, Beyanname No, Miktar, Tutar")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--gun", type=int, default=60, help="Süre sonuna ve kapatma son gününe kaç gün kala uyarılsın (varsayılan 60)")
    p.add_argument("--kapatma-ay", type=int, default=3, help="Kapatma başvurusu için süre sonundan itibaren tanınan ay (varsayılan 3)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "diib_takibi.xlsx")
    a = p.parse_args(argv)
    if (a.belgeler is None) != (a.hareketler is None):
        print("[X] --belgeler ve --hareketler birlikte verilmelidir")
        return 2
    ornek = a.belgeler is None
    b_yol, h_yol = (ORNEK / "belge_satirlari.csv", ORNEK / "hareketler.csv") if ornek else (a.belgeler, a.hareketler)
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    if bugun is None or not 0 <= a.gun <= 365 or not 0 <= a.kapatma_ay <= 24:
        print("[X] --tarih GG.AA.YYYY; --gun 0–365; --kapatma-ay 0–24 arasında olmalı")
        return 2
    for y in (b_yol, h_yol):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(b_yol, h_yol, a.cikti, bugun, a.gun, a.kapatma_ay)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(r['belgeler'])} belge · " + ", ".join(f"{d}: {n}" for d, n in r["say"].items() if n))
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
