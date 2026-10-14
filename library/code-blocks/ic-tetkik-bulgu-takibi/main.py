"""
İç Tetkik Bulgu Takibi — Workers / Workless kod bloğu
Kalite Güvence ve Kalite Kontrol › Kalite Müdürü
ISO 9001 iç tetkik bulgularını madde, birim ve kapanış durumuna göre takip eder; geciken, etkinliği doğrulanmamış ve tekrar eden bulguları listeler.
  Termin        = satırdaki termin; yoksa tetkik tarihi + türe göre hedef süre (majör --major 30, minör --minor 60 gün; gözlem ve iyileştirme fırsatında termin aranmaz)
  Gecikme       = açık bulguda bugün − termin
  Kapanış süresi = kapanış tarihi − tetkik tarihi (gün)
  Tekrar eden   = aynı birim + madde, en az iki farklı tetkikte uygunsuzluk
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
SUTUN = {"no": ("bulgu no", "no", "bulgu"), "tetkik": ("tetkik no", "tetkik", "denetim no"), "tetkik_tarih": ("tetkik tarihi", "denetim tarihi", "tarih"), "birim": ("birim", "surec", "bolum"),
         "madde": ("madde", "standart maddesi", "iso madde", "madde no"), "tur": ("tur", "bulgu turu", "siniflandirma"), "aciklama": ("aciklama", "bulgu aciklamasi", "tespit"),
         "sorumlu": ("sorumlu", "faaliyet sorumlusu"), "termin": ("termin", "hedef tarih", "kapanis hedefi"), "kok": ("kok neden", "kok neden analizi"),
         "dof": ("duzeltici faaliyet", "dof", "faaliyet"), "kapanis": ("kapanis tarihi", "kapanis", "kapatma tarihi"), "etkinlik": ("etkinlik dogrulama", "etkinlik", "etkinlik degerlendirmesi")}
TURLER = ("Majör uygunsuzluk", "Minör uygunsuzluk", "Gözlem", "İyileştirme fırsatı")


def _tur(x) -> str | None:
    k = katla(x)
    if "major" in k or k in ("buyuk", "buyuk uygunsuzluk"):
        return "Majör uygunsuzluk"
    if "minor" in k or k in ("kucuk", "kucuk uygunsuzluk", "uygunsuzluk"):
        return "Minör uygunsuzluk"
    if "gozlem" in k:
        return "Gözlem"
    if "iyilestirme" in k or "firsat" in k or "oneri" in k:
        return "İyileştirme fırsatı"
    return None


def _etkinlik(x) -> str:
    k = katla(x)
    if k in ("etkin", "etkili", "evet", "uygun"):
        return "Etkin"
    if "degil" in k or k in ("hayir", "etkisiz", "yetersiz"):
        return "Etkin değil"
    return "Bekliyor" if not k or "bekl" in k else metin(x)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("birim", "tur")):
        if not metin(r["birim"]) and not metin(r["no"]):
            continue
        no = metin(r["no"]) or f"Satır {r['_satir']}"
        tur = _tur(r["tur"])
        t = {a: (tarih(r[a]) if metin(r[a]) else None) for a in ("tetkik_tarih", "termin", "kapanis")}
        if tur is None or t["tetkik_tarih"] is None or any(metin(r[a]) and v is None for a, v in t.items()):
            uy.append(uyari("Yüksek", "Satır okunamadı", no, f"Satır {r['_satir']}: tür ('{metin(r['tur'])}'), tetkik tarihi, termin ya da kapanış tarihi okunamadı."))
            continue
        if katla(no) in gorulen:
            uy.append(uyari("Orta", "Bulgu iki kez yazılmış", no, f"Satır {r['_satir']}: ikinci satır alınmadı."))
            continue
        gorulen.add(katla(no))
        l.append({"no": no, "tetkik": metin(r["tetkik"]) or f"{t['tetkik_tarih']:%Y-%m}", "birim": metin(r["birim"]), "madde": metin(r["madde"]) or "-", "tur": tur,
                  "aciklama": metin(r["aciklama"]), "sorumlu": metin(r["sorumlu"]), "kok": metin(r["kok"]), "dof": metin(r["dof"]), "etkinlik": _etkinlik(r["etkinlik"]), **t})
    return l, uy


def analiz(bulgular: list[dict], bugun: date, major: int = 30, minor: int = 60) -> dict:
    if not bulgular:
        raise ValueError("Okunabilen bulgu yok.")
    uy = []
    for b in bulgular:
        kim = f"{b['no']} · {b['birim']} · {b['madde']}"
        uygunsuz = b["tur"] in TURLER[:2]
        b["hedef"] = b["termin"] or (b["tetkik_tarih"] + timedelta(days=major if b["tur"] == TURLER[0] else minor) if uygunsuz else None)
        b["acik"] = b["kapanis"] is None
        b["gecikme"] = (bugun - b["hedef"]).days if b["acik"] and b["hedef"] and bugun > b["hedef"] else None
        b["sure"] = (b["kapanis"] - b["tetkik_tarih"]).days if b["kapanis"] else None
        if b["kapanis"] and b["kapanis"] < b["tetkik_tarih"]:
            uy.append(uyari("Orta", "Kapanış tetkikten önce", kim, f"Tetkik {b['tetkik_tarih']:%d.%m.%Y}, kapanış {b['kapanis']:%d.%m.%Y}."))
        if b["gecikme"]:
            uy.append(uyari("Yüksek" if b["tur"] == TURLER[0] else "Orta", "Termin geçti", kim, f"{b['tur']}: termin {b['hedef']:%d.%m.%Y}, {b['gecikme']} gün gecikti." + (f" Sorumlu: {b['sorumlu']}." if b["sorumlu"] else " Sorumlu yazılmamış.")))
        if uygunsuz and b["acik"] and not b["sorumlu"]:
            uy.append(uyari("Orta", "Sorumlu yok", kim, "Açık uygunsuzluğun sorumlusu yazılmamış."))
        if uygunsuz and not b["acik"]:
            if not b["kok"] or not b["dof"]:
                uy.append(uyari("Orta", "Kök neden / faaliyet eksik", kim, "Kapatılmış uygunsuzlukta " + " ve ".join(a for a, v in (("kök neden", b["kok"]), ("düzeltici faaliyet", b["dof"])) if not v) + " yazılmamış."))
            if b["etkinlik"] == "Etkin değil":
                uy.append(uyari("Orta", "Faaliyet etkin değil", kim, "Etkinlik doğrulaması olumsuz; bulgu yeniden açılmalı ya da yeni faaliyet planlanmalı."))
            elif b["etkinlik"] == "Bekliyor":
                uy.append(uyari("Bilgi", "Etkinlik doğrulanmadı", kim, f"Kapanış {b['kapanis']:%d.%m.%Y}; düzeltici faaliyetin etkinliği henüz doğrulanmamış."))
    gruplar = {}
    for b in bulgular:
        if b["tur"] in TURLER[:2]:
            gruplar.setdefault((katla(b["birim"]), katla(b["madde"])), []).append(b)
    tekrar = []
    for l in gruplar.values():
        tetkikler = {b["tetkik"] for b in l}
        if len(tetkikler) >= 2:
            tekrar.append({"birim": l[0]["birim"], "madde": l[0]["madde"], "tetkik": len(tetkikler), "bulgular": [b["no"] for b in sorted(l, key=lambda b: b["tetkik_tarih"])]})
            uy.append(uyari("Orta", "Tekrar eden bulgu", f"{l[0]['birim']} · {l[0]['madde']}", f"{len(tetkikler)} farklı tetkikte uygunsuzluk ({', '.join(b['no'] for b in l)}); önceki faaliyetlerin etkinliğini sorgulayın."))

    def ozet(fn, sira=None):
        g = {}
        for b in bulgular:
            g.setdefault(fn(b), []).append(b)
        out = []
        for ad, l in g.items():
            sure = [b["sure"] for b in l if b["sure"] is not None and b["sure"] >= 0]
            out.append({"ad": ad, "bulgu": len(l), "major": sum(1 for b in l if b["tur"] == TURLER[0]), "minor": sum(1 for b in l if b["tur"] == TURLER[1]),
                        "diger": sum(1 for b in l if b["tur"] in TURLER[2:]), "acik": sum(1 for b in l if b["acik"]), "geciken": sum(1 for b in l if b["gecikme"]),
                        "ort": yuvarla(Decimal(sum(sure)) / len(sure), "0.1") if sure else None, "etkin_degil": sum(1 for b in l if b["etkinlik"] == "Etkin değil")})
        if sira:
            return sorted(out, key=lambda s: sira(s["ad"]))
        return sorted(out, key=lambda s: (-s["geciken"], -s["acik"], -s["bulgu"], katla(s["ad"])))

    def madde_sira(m):
        return tuple(int(x) if x.isdigit() else 999 for x in re.split(r"[.\s]", m) if x) or (999,)
    acik = sorted((b for b in bulgular if b["acik"]), key=lambda b: (TURLER.index(b["tur"]), -(b["gecikme"] or 0), b["hedef"] or date.max))
    return {"acik": acik, "bulgular": bulgular, "birim": ozet(lambda b: b["birim"]), "madde": ozet(lambda b: b["madde"], madde_sira), "tetkik": ozet(lambda b: b["tetkik"], lambda x: x),
            "toplam": ozet(lambda b: "TOPLAM")[0], "tekrar": tekrar, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Açık Bulgular"
    _baslik(ws, ["Bulgu No", "Tetkik", "Tetkik Tarihi", "Birim", "Madde", "Tür", "Açıklama", "Sorumlu", "Termin", "Gecikme (gün)"], (10, 10, 11, 16, 8, 18, 44, 14, 11, 9), "B2")
    for b in r["acik"]:
        _satir(ws, [b["no"], b["tetkik"], b["tetkik_tarih"], b["birim"], b["madde"], b["tur"], b["aciklama"], b["sorumlu"], b["hedef"], b["gecikme"]], tarih_s=(3, 9), sar=(7,),
               dolgu={10: KIRMIZI if b["gecikme"] and b["tur"] == TURLER[0] else (SARI if b["gecikme"] else None)})
    bas = ["Bulgu", "Majör", "Minör", "Gözlem / İyileştirme", "Açık", "Geciken", "Ort. Kapanış (gün)", "Etkin Değil"]
    for ad, l in (("Birim", r["birim"]), ("Madde", r["madde"]), ("Tetkik", r["tetkik"])):
        ws = wb.create_sheet(ad)
        _baslik(ws, [ad] + bas, (18, 7, 7, 7, 10, 7, 8, 10, 9))
        for s in l + ([r["toplam"]] if ad == "Birim" else []):
            _satir(ws, [s["ad"], s["bulgu"], s["major"], s["minor"], s["diger"], s["acik"], s["geciken"], s["ort"], s["etkin_degil"]], dolgu={7: SARI if s["geciken"] else None})
    ws = wb.create_sheet("Tekrar Eden")
    _baslik(ws, ["Birim", "Madde", "Tetkik Sayısı", "Bulgular"], (18, 8, 10, 40))
    for t in r["tekrar"]:
        _satir(ws, [t["birim"], t["madde"], t["tetkik"], ", ".join(t["bulgular"])])
    ws = wb.create_sheet("Tüm Bulgular")
    _baslik(ws, ["Bulgu No", "Tetkik", "Tetkik Tarihi", "Birim", "Madde", "Tür", "Sorumlu", "Termin", "Kapanış", "Kapanış Süresi (gün)", "Kök Neden", "Düzeltici Faaliyet", "Etkinlik"],
            (10, 10, 11, 16, 8, 18, 14, 11, 11, 9, 30, 30, 11))
    for b in sorted(r["bulgular"], key=lambda b: (b["tetkik_tarih"], b["no"])):
        _satir(ws, [b["no"], b["tetkik"], b["tetkik_tarih"], b["birim"], b["madde"], b["tur"], b["sorumlu"], b["hedef"], b["kapanis"], b["sure"], b["kok"], b["dof"], b["etkinlik"]],
               tarih_s=(3, 8, 9), sar=(11, 12))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, major: int = 30, minor: int = 60) -> dict:
    b, uy = oku(yol)
    r = analiz(b, bugun, major, minor)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="İç tetkik bulgularını madde, birim ve kapanış durumuna göre takip eder.")
    p.add_argument("--bulgular", type=Path, help="Bulgu No, Tetkik No, Tetkik Tarihi, Birim, Madde, Tür (Majör / Minör / Gözlem / İyileştirme fırsatı), Açıklama, Sorumlu, Termin, "
                                                 "Kök Neden, Düzeltici Faaliyet, Kapanış Tarihi, Etkinlik Doğrulama")
    p.add_argument("--tarih", help="Rapor günü GG.AA.YYYY (varsayılan bugün; örnek veride 14.10.2026)")
    p.add_argument("--major", type=int, default=30, help="Termin yazılmamış majör uygunsuzluk için hedef kapanış süresi, gün (varsayılan 30)")
    p.add_argument("--minor", type=int, default=60, help="Termin yazılmamış minör uygunsuzluk için hedef kapanış süresi, gün (varsayılan 60)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "ic_tetkik_bulgulari.xlsx")
    a = p.parse_args(argv)
    yol = a.bulgular or ORNEK / "bulgular.csv"
    bugun = tarih(a.tarih) if a.tarih else (date(2026, 10, 14) if a.bulgular is None else date.today())
    if bugun is None or a.major < 1 or a.minor < 1:
        print("[X] --tarih GG.AA.YYYY; --major ve --minor en az 1 olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, bugun, a.major, a.minor)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {t['bulgu']} bulgu (majör {t['major']}, minör {t['minor']}) · açık {t['acik']} · geciken {t['geciken']} · ort. kapanış {t['ort']} gün · tekrar eden {len(r['tekrar'])}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
