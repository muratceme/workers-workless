import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 14)


def b(no, tetkik, tt, birim, madde, tur, termin=None, kapanis=None, kok="K", dof="D", etkinlik="Bekliyor", sorumlu="S"):
    return {"no": no, "tetkik": tetkik, "tetkik_tarih": tt, "birim": birim, "madde": madde, "tur": tur, "aciklama": "", "sorumlu": sorumlu, "kok": kok, "dof": dof, "etkinlik": etkinlik,
            "termin": termin, "kapanis": kapanis}


class AnalizTesti(unittest.TestCase):
    def test_analiz(self):
        t1, t2 = date(2026, 5, 1), date(2026, 9, 1)
        l = [b("1", "A", t1, "Üretim", "7.5.3", "Majör uygunsuzluk"), b("2", "A", t1, "Depo", "8.7", "Minör uygunsuzluk", kapanis=date(2026, 6, 10), kok="", etkinlik="Etkin değil"),
             b("3", "B", t2, "Üretim", "7.5.3", "Minör uygunsuzluk", termin=date(2026, 10, 20), sorumlu=""), b("4", "B", t2, "Depo", "10.2", "Gözlem"),
             b("5", "B", t2, "Satış", "9.1.2", "Minör uygunsuzluk", kapanis=date(2026, 9, 21))]
        r = main.analiz(l, BUGUN)
        x = {v["no"]: v for v in r["bulgular"]}
        self.assertEqual((x["1"]["hedef"], x["1"]["gecikme"], x["3"]["hedef"], x["3"]["gecikme"], x["4"]["hedef"], x["2"]["sure"]), (date(2026, 5, 31), 136, date(2026, 10, 20), None, None, 40))
        self.assertEqual([v["no"] for v in r["acik"]], ["1", "3", "4"])
        self.assertEqual(r["tekrar"], [{"birim": "Üretim", "madde": "7.5.3", "tetkik": 2, "bulgular": ["1", "3"]}])
        t = r["toplam"]
        self.assertEqual((t["bulgu"], t["major"], t["minor"], t["diger"], t["acik"], t["geciken"], t["ort"], t["etkin_degil"]), (5, 1, 3, 1, 3, 1, D("30.0"), 1))
        self.assertEqual([s["ad"] for s in r["madde"]], ["7.5.3", "8.7", "9.1.2", "10.2"])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"].split(" ·")[0]) for u in r["uyarilar"]),
                         [("Bilgi", "Etkinlik doğrulanmadı", "5"), ("Orta", "Faaliyet etkin değil", "2"), ("Orta", "Kök neden / faaliyet eksik", "2"), ("Orta", "Sorumlu yok", "3"),
                          ("Orta", "Tekrar eden bulgu", "Üretim"), ("Yüksek", "Termin geçti", "1")])
        self.assertEqual(main.analiz(l, BUGUN, major=200)["bulgular"][0]["gecikme"], None)
        with self.assertRaises(ValueError):
            main.analiz([], BUGUN)

    def test_donusum(self):
        self.assertEqual([main._tur(x) for x in ("Majör", "minor uygunsuzluk", "Uygunsuzluk", "GÖZLEM", "İyileştirme önerisi", "x")],
                         ["Majör uygunsuzluk", "Minör uygunsuzluk", "Minör uygunsuzluk", "Gözlem", "İyileştirme fırsatı", None])
        self.assertEqual([main._etkinlik(x) for x in ("Etkin", "Etkin değil", "", "beklemede")], ["Etkin", "Etkin değil", "Bekliyor", "Bekliyor"])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "bulgular.csv", cikti, BUGUN)
            t = r["toplam"]
            self.assertEqual((t["bulgu"], t["major"], t["minor"], t["acik"], t["geciken"], t["ort"], len(r["tekrar"])), (35, 3, 22, 8, 3, D("38.4"), 5))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Açık Bulgular", "Birim", "Madde", "Tetkik", "Tekrar Eden", "Tüm Bulgular", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "b.csv"
            y.write_text("Bulgu;Tarih;Süreç;Madde No;Sınıflandırma;Kapanış\n1;01.09.2026;Üretim;8.5.1;Minör;\n1;01.09.2026;Üretim;8.5.1;Minör;\n2;x;Depo;8.7;Minör;\n3;01.09.2026;Depo;8.7;belki;\n"
                         "4;01.09.2026;Depo;8.7;Gözlem;x\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["tetkik"]) for x in l], [u["tur"] for u in uy]), ([("1", "2026-09")], ["Bulgu iki kez yazılmış", "Satır okunamadı", "Satır okunamadı", "Satır okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--bulgular", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "14.10.2026", "--major", "15", "--minor", "45"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--minor", "0"])], [2, 2])
            self.assertEqual(main.main(["--bulgular", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Süreç;Sınıflandırma\nA;x\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
