"""
Kanallar Arası Stok ve Fiyat Kontrolü — Workers / Workless kod bloğu
Pazaryeri Yönetimi › Pazaryeri Uzmanı

Farklı satış kanallarındaki stok ve fiyatları merkezi listeyle karşılaştırıp uyumsuzlukları listeler.
  Merkezde yok              = kanalda satışta olan stok kodu merkezi listede bulunmuyor
  Stok yokken satışta       = merkez stok 0, kanal stok > 0                       (olmayan ürün satılır: iptal ve ceza riski)
  Kanal stoku merkezden fazla = kanal stok − merkez stok > --stok-tolerans (0)
  Kanalda stok kapalı       = merkez stok > 0, kanal stok 0                       (satış kaybı)
  Kanal stoku eksik         = merkez stok − kanal stok > --stok-tolerans; yalnız --tam-stok verilirse (kanala stokun tamamını açıyorsanız)
  Fiyat düşük / yüksek      = (kanal fiyatı − beklenen fiyat) ÷ beklenen fiyat, mutlak değeri --fiyat-tolerans (%1) üzerinde
  Beklenen fiyat            = merkez fiyat × (1 + kanal fiyat farkı %)             (kanal ayar dosyası verilmezse fark 0)
  Kanalda olmayan ürün      = merkezde stoku olan ama o kanalda hiç listelenmemiş stok kodu (ayrı sayfada, bilgi amaçlı)
Kanal panellerine bağlanmaz; panellerden indirdiğiniz listeleri karşılaştırır. Listeler aynı anda alınmış olmalıdır.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


MERKEZ = {"sku": ("stok kodu", "sku", "barkod", "urun kodu"), "urun": ("urun", "urun adi"), "stok": ("stok", "stok adedi", "mevcut stok"), "fiyat": ("fiyat", "satis fiyati", "liste fiyati")}
KANAL = {"kanal": ("kanal", "pazaryeri", "satis kanali"), "sku": MERKEZ["sku"], "urun": MERKEZ["urun"], "stok": MERKEZ["stok"], "fiyat": MERKEZ["fiyat"]}
AYAR = {"kanal": KANAL["kanal"], "fark": ("fiyat farki", "fiyat farki yuzde", "fark", "fiyat farki %")}
ONEM = {"Merkezde yok": "Yüksek", "Stok yokken satışta": "Yüksek", "Fiyat düşük": "Yüksek", "Kanal stoku merkezden fazla": "Orta", "Kanalda stok kapalı": "Orta", "Fiyat yüksek": "Orta",
        "Kanal stoku eksik": "Bilgi"}
Z = Decimal(0)


def oku(merkez_yol: Path, kanal_yol: Path, ayar_yol: Path | None = None) -> tuple[dict, list[dict], dict, list[dict]]:
    uy, merkez, kanal, ayar = [], {}, [], {}
    for r in kayitlar(merkez_yol, MERKEZ, ("sku", "stok")):
        s = metin(r["sku"])
        if not s:
            continue
        kim = f"{s} ({merkez_yol.name} satır {r['_satir']})"
        stok, fiyat = para(r["stok"]), para(r["fiyat"])
        if stok is None:
            uy.append(uyari("Yüksek", "Merkez stoku okunamadı", kim, f"Stok '{metin(r['stok'])}'. Ürün merkezi listeye alınmadı."))
            continue
        if katla(s) in merkez:
            uy.append(uyari("Yüksek", "Mükerrer stok kodu", kim, "Merkezi listede aynı stok kodu birden fazla; yalnız ilki alındı."))
            continue
        if fiyat is None and metin(r["fiyat"]):
            uy.append(uyari("Orta", "Merkez fiyatı okunamadı", kim, f"Fiyat '{metin(r['fiyat'])}'. Bu ürün için fiyat kontrolü yapılmadı."))
        merkez[katla(s)] = {"sku": s, "urun": metin(r["urun"]), "stok": max(stok, Z), "fiyat": fiyat if fiyat and fiyat > 0 else None}
    gorulen = set()
    for r in kayitlar(kanal_yol, KANAL, ("kanal", "sku")):
        k, s = metin(r["kanal"]), metin(r["sku"])
        if not k or not s:
            continue
        kim = f"{k} · {s} ({kanal_yol.name} satır {r['_satir']})"
        if (katla(k), katla(s)) in gorulen:
            uy.append(uyari("Yüksek", "Mükerrer kanal kaydı", kim, "Aynı kanalda aynı stok kodu birden fazla; yalnız ilki alındı. Kanalda çift ilan olabilir."))
            continue
        gorulen.add((katla(k), katla(s)))
        stok, fiyat = para(r["stok"]), para(r["fiyat"])
        for v, ad, ham in ((stok, "Stok", r["stok"]), (fiyat, "Fiyat", r["fiyat"])):
            if v is None and metin(ham):
                uy.append(uyari("Orta", "Kanal değeri okunamadı", kim, f"{ad} '{metin(ham)}'. Bu alan için kontrol yapılmadı."))
        kanal.append({"kanal": k, "sku": s, "urun": metin(r["urun"]), "stok": stok, "fiyat": fiyat})
    if ayar_yol:
        for r in kayitlar(ayar_yol, AYAR, ("kanal", "fark")):
            k, f = metin(r["kanal"]), para(r["fark"])
            if not k:
                continue
            if f is None or f <= -100:
                uy.append(uyari("Yüksek", "Kanal fiyat farkı okunamadı", k, f"{ayar_yol.name} satır {r['_satir']}: '{metin(r['fark'])}'. Bu kanal için fark 0 alındı."))
                continue
            ayar[katla(k)] = f
    return merkez, kanal, ayar, uy


def karsilastir(merkez: dict, kanal: list[dict], ayar: dict | None = None, fiyat_tolerans: Decimal = Decimal(1), stok_tolerans: Decimal = Z, tam_stok: bool = False) -> dict:
    if not merkez:
        raise ValueError("Merkezi listede ürün yok.")
    if not kanal:
        raise ValueError("Kanal listesinde kayıt yok.")
    ayar = ayar or {}
    for x in kanal:
        m, s = merkez.get(katla(x["sku"])), []
        x["merkez_stok"] = x["merkez_fiyat"] = x["beklenen"] = x["fiyat_farki"] = x["stok_farki"] = None
        if m is None:
            s.append("Merkezde yok")
        else:
            x["urun"] = x["urun"] or m["urun"]
            x["merkez_stok"], x["merkez_fiyat"] = m["stok"], m["fiyat"]
            if x["stok"] is not None:
                x["stok_farki"] = x["stok"] - m["stok"]
                if m["stok"] == 0 and x["stok"] > 0:
                    s.append("Stok yokken satışta")
                elif x["stok_farki"] > stok_tolerans:
                    s.append("Kanal stoku merkezden fazla")
                elif m["stok"] > 0 and x["stok"] <= 0:
                    s.append("Kanalda stok kapalı")
                elif tam_stok and -x["stok_farki"] > stok_tolerans:
                    s.append("Kanal stoku eksik")
            if x["fiyat"] is not None and m["fiyat"]:
                x["beklenen"] = yuvarla(m["fiyat"] * (1 + ayar.get(katla(x["kanal"]), Z) / 100))
                x["fiyat_farki"] = yuvarla((x["fiyat"] - x["beklenen"]) / x["beklenen"] * 100)
                if abs(x["fiyat_farki"]) > fiyat_tolerans:
                    s.append("Fiyat düşük" if x["fiyat_farki"] < 0 else "Fiyat yüksek")
        x["sorun"] = s
        x["onem"] = min((ONEM[t] for t in s), key=("Yüksek", "Orta", "Bilgi").index) if s else ""
    kanallar = {}
    for x in kanal:
        k = kanallar.setdefault(katla(x["kanal"]), {"kanal": x["kanal"], "kayit": 0, "uyumlu": 0, "fark": ayar.get(katla(x["kanal"]), Z), "skular": set()} | {t: 0 for t in ONEM})
        k["kayit"] += 1
        k["uyumlu"] += not x["sorun"]
        k["skular"].add(katla(x["sku"]))
        for t in x["sorun"]:
            k[t] += 1
    olmayan = [{"kanal": k["kanal"], **m} for k in kanallar.values() for a, m in merkez.items() if m["stok"] > 0 and a not in k["skular"]]
    for k in kanallar.values():
        k["olmayan"] = sum(o["kanal"] == k["kanal"] for o in olmayan)
    sira = ("Yüksek", "Orta", "Bilgi", "")
    return {"kayitlar": sorted(kanal, key=lambda x: (sira.index(x["onem"]), x["kanal"], x["sku"])), "kanallar": sorted(kanallar.values(), key=lambda k: k["kanal"]),
            "olmayan": sorted(olmayan, key=lambda o: (o["kanal"], o["sku"])), "ayar": (fiyat_tolerans, stok_tolerans, tam_stok),
            "say": {"kayit": len(kanal), "sorunlu": sum(bool(x["sorun"]) for x in kanal), "yuksek": sum(x["onem"] == "Yüksek" for x in kanal)},
            "eksik_ayar": sorted(k for k in ayar if k not in kanallar)}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    ft, st, tam = r["ayar"]
    ws = wb.active
    ws.title = "Uyumsuzluklar"
    _baslik(ws, ["Önem", "Kanal", "Stok Kodu", "Ürün", "Sorun", "Merkez Stok", "Kanal Stok", "Stok Farkı", "Merkez Fiyat", "Beklenen Kanal Fiyatı", "Kanal Fiyatı", "Fiyat Farkı %", "Yapılan (siz doldurun)"],
            (9, 20, 16, 32, 34, 10, 10, 9, 12, 12, 12, 10, 26))
    for x in r["kayitlar"]:
        if x["sorun"]:
            _satir(ws, [x["onem"], x["kanal"], x["sku"], x["urun"], ", ".join(x["sorun"]), x["merkez_stok"], x["stok"], x["stok_farki"], x["merkez_fiyat"], x["beklenen"], x["fiyat"], x["fiyat_farki"], ""],
                   para_s=(9, 10, 11, 12), dolgu={1: RENK[x["onem"]]}, sar=(5,))
    ws.append([])
    ws.append([f"Fiyat toleransı %{ft.normalize():f}, stok toleransı {st.normalize():f} adet{'; kanala stokun tamamı açılıyor kabul edildi' if tam else ''}. Listeler aynı anda alınmadıysa farkların bir kısmı zamanlamadan kaynaklanır."])
    ws = wb.create_sheet("Kanal Özeti")
    turler = list(ONEM)
    _baslik(ws, ["Kanal", "Fiyat Farkı %", "Kayıt", "Uyumlu", *turler, "Kanalda Olmayan Ürün"], (22, 10, 8, 8, 10, 12, 9, 14, 12, 9, 10, 12))
    for k in r["kanallar"]:
        _satir(ws, [k["kanal"], k["fark"], k["kayit"], k["uyumlu"], *[k[t] or None for t in turler], k["olmayan"] or None], para_s=(2,))
    ws = wb.create_sheet("Kanalda Olmayan Ürünler")
    _baslik(ws, ["Kanal", "Stok Kodu", "Ürün", "Merkez Stok", "Merkez Fiyat"], (22, 16, 36, 10, 12))
    for o in r["olmayan"]:
        _satir(ws, [o["kanal"], o["sku"], o["urun"], o["stok"], o["fiyat"]], para_s=(5,))
    ws = wb.create_sheet("Tüm Kayıtlar")
    _baslik(ws, ["Kanal", "Stok Kodu", "Ürün", "Durum", "Merkez Stok", "Kanal Stok", "Merkez Fiyat", "Beklenen Kanal Fiyatı", "Kanal Fiyatı", "Fiyat Farkı %"], (20, 16, 32, 34, 10, 10, 12, 12, 12, 10))
    for x in sorted(r["kayitlar"], key=lambda x: (x["kanal"], x["sku"])):
        _satir(ws, [x["kanal"], x["sku"], x["urun"], ", ".join(x["sorun"]) or "Uyumlu", x["merkez_stok"], x["stok"], x["merkez_fiyat"], x["beklenen"], x["fiyat"], x["fiyat_farki"]], para_s=(7, 8, 9, 10),
               dolgu={4: RENK[x["onem"]] if x["onem"] else None})
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(merkez_yol: Path, kanal_yol: Path, cikti: Path, ayar_yol: Path | None = None, **ayar) -> dict:
    merkez, kanal, kanal_ayar, uy = oku(merkez_yol, kanal_yol, ayar_yol)
    r = karsilastir(merkez, kanal, kanal_ayar, **ayar)
    for k in r["eksik_ayar"]:
        uy.append(uyari("Orta", "Ayarı olan kanalın kaydı yok", k, "Kanal ayar dosyasında var, kanal listesinde bu adla kayıt yok. Kanal adlarının iki dosyada aynı yazıldığını kontrol edin."))
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Kanal stok ve fiyatlarını merkezi listeyle karşılaştırır.")
    p.add_argument("--merkez", type=Path, default=ORNEK / "merkez_liste.csv", help="Stok Kodu, Ürün, Stok, Fiyat")
    p.add_argument("--kanal", type=Path, default=ORNEK / "kanal_listesi.csv", help="Kanal, Stok Kodu, Stok, Fiyat (tüm kanallar alt alta)")
    p.add_argument("--kanal-ayar", type=Path, help="İsteğe bağlı: Kanal, Fiyat Farkı %% (kanalda bilerek uyguladığınız fiyat farkı)")
    p.add_argument("--fiyat-tolerans", default="1", help="Bu %%'nin üzerindeki fiyat farkı uyumsuzluk sayılır (varsayılan 1)")
    p.add_argument("--stok-tolerans", default="0", help="Bu adede kadar stok farkı görmezden gelinir (varsayılan 0)")
    p.add_argument("--tam-stok", action="store_true", help="Kanala stokun tamamını açıyorsanız: kanal stoku merkezden azsa da listeler")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kanal_stok_fiyat_kontrolu.xlsx")
    a = p.parse_args(argv)
    ayar_yol = a.kanal_ayar or (ORNEK / "kanal_ayarlari.csv" if a.kanal == ORNEK / "kanal_listesi.csv" else None)
    ft, st = para(a.fiyat_tolerans), para(a.stok_tolerans)
    if ft is None or st is None or ft < 0 or st < 0:
        print("[X] --fiyat-tolerans ve --stok-tolerans sıfır veya pozitif sayı olmalı")
        return 2
    for y in (a.merkez, a.kanal, ayar_yol):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.merkez, a.kanal, a.cikti, ayar_yol, fiyat_tolerans=ft, stok_tolerans=st, tam_stok=a.tam_stok)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    s = r["say"]
    print(f"[OK] {len(r['kanallar'])} kanal · {s['kayit']} kayıt · {s['sorunlu']} uyumsuz ({s['yuksek']} yüksek önemli) · {len(r['olmayan'])} kanalda olmayan ürün")
    for x in r["kayitlar"]:
        if x["onem"] == "Yüksek":
            print(f"[!] {x['kanal']} · {x['sku']}: {', '.join(x['sorun'])}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
