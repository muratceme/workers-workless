import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def mrk(**u):
    return {k.lower(): {"sku": k, "urun": k, "stok": D(s), "fiyat": None if f is None else D(f)} for k, (s, f) in u.items()}


def kn(sku, stok, fiyat, kanal="K"):
    return {"kanal": kanal, "sku": sku, "urun": "", "stok": None if stok is None else D(stok), "fiyat": None if fiyat is None else D(fiyat)}


def tek(mstok, mfiyat, kstok, kfiyat, **kw):
    return main.karsilastir(mrk(A=(mstok, mfiyat)), [kn("A", kstok, kfiyat)], **kw)["kayitlar"][0]


class KarsilastirmaTesti(unittest.TestCase):
    def test_stok_kurallari(self):
        self.assertEqual([tek(m, 100, k, 100)["sorun"] for m, k in ((10, 10), (0, 0), (0, 3), (5, 9), (5, 0), (10, 4))],
                         [[], [], ["Stok yokken satışta"], ["Kanal stoku merkezden fazla"], ["Kanalda stok kapalı"], []])
        self.assertEqual(tek(10, 100, 4, 100, tam_stok=True)["sorun"], ["Kanal stoku eksik"])
        self.assertEqual([tek(10, 100, k, 100, stok_tolerans=D(2), tam_stok=True)["sorun"] for k in (12, 13, 8, 7)], [[], ["Kanal stoku merkezden fazla"], [], ["Kanal stoku eksik"]])
        x = tek(0, 100, 3, 100)
        self.assertEqual((x["onem"], x["stok_farki"]), ("Yüksek", D(3)))
        self.assertEqual(tek(5, 100, None, 100)["sorun"], [])

    def test_fiyat_kurallari(self):
        self.assertEqual([tek(5, 100, 5, f)["sorun"] for f in (101, "101.01", 99, "98.99")], [[], ["Fiyat yüksek"], [], ["Fiyat düşük"]])            # tam %1 uyumlu
        x = main.karsilastir(mrk(A=(8, 2000)), [kn("A", 8, 2000)], {"k": D(10)})["kayitlar"][0]
        self.assertEqual((x["beklenen"], x["fiyat_farki"], x["sorun"], x["onem"]), (D("2200.00"), D("-9.09"), ["Fiyat düşük"], "Yüksek"))
        self.assertEqual(tek(5, 100, 5, 110, fiyat_tolerans=D(10))["sorun"], [])
        self.assertEqual((tek(5, None, 5, 110)["sorun"], tek(5, 100, 5, None)["fiyat_farki"]), ([], None))
        x = tek(0, 100, 2, 50)
        self.assertEqual((x["sorun"], x["onem"]), (["Stok yokken satışta", "Fiyat düşük"], "Yüksek"))

    def test_merkezde_yok_olmayan_ve_ozet(self):
        r = main.karsilastir(mrk(A=(5, 100), B=(0, 100), C=(3, 100)), [kn("A", 5, 100), kn("Z", 1, 10), kn("a", 5, 120, "L")], {"m": D(5)})
        self.assertEqual([(x["kanal"], x["sku"], x["sorun"]) for x in r["kayitlar"]], [("K", "Z", ["Merkezde yok"]), ("L", "a", ["Fiyat yüksek"]), ("K", "A", [])])
        self.assertEqual([(o["kanal"], o["sku"]) for o in r["olmayan"]], [("K", "C"), ("L", "C")])                    # stoku olmayan B listelenmez
        self.assertEqual([(k["kanal"], k["kayit"], k["uyumlu"], k["Merkezde yok"], k["Fiyat yüksek"], k["olmayan"]) for k in r["kanallar"]], [("K", 2, 1, 1, 0, 1), ("L", 1, 0, 0, 1, 1)])
        self.assertEqual((r["say"], r["eksik_ayar"]), ({"kayit": 3, "sorunlu": 2, "yuksek": 1}, ["m"]))
        for m, k in (({}, [kn("A", 1, 1)]), (mrk(A=(1, 1)), [])):
            with self.assertRaises(ValueError):
                main.karsilastir(m, k)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "k.xlsx"
            r = main.calistir(ORNEK / "merkez_liste.csv", ORNEK / "kanal_listesi.csv", cikti, ORNEK / "kanal_ayarlari.csv")
            self.assertEqual({(x["kanal"], x["sku"]): x["sorun"] for x in r["kayitlar"] if x["sorun"]},
                             {("Web Sitesi", "SKU-004"): ["Fiyat düşük"], ("Pazaryeri A", "SKU-002"): ["Stok yokken satışta"], ("Pazaryeri A", "SKU-003"): ["Kanal stoku merkezden fazla"],
                              ("Pazaryeri A", "SKU-004"): ["Fiyat düşük"], ("Pazaryeri A", "SKU-009"): ["Merkezde yok"], ("Pazaryeri B", "SKU-001"): ["Kanalda stok kapalı"],
                              ("Pazaryeri B", "SKU-003"): ["Fiyat yüksek"]})
            self.assertEqual((r["say"], [(o["kanal"], o["sku"]) for o in r["olmayan"]], r["uyarilar"]),
                             ({"kayit": 13, "sorunlu": 7, "yuksek": 4}, [("Pazaryeri A", "SKU-005"), ("Pazaryeri B", "SKU-004")], []))
            self.assertEqual(next(x["fiyat_farki"] for x in r["kayitlar"] if (x["kanal"], x["sku"]) == ("Pazaryeri B", "SKU-005")), D("0.67"))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Uyumsuzluklar", "Kanal Özeti", "Kanalda Olmayan Ürünler", "Tüm Kayıtlar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            m, k, ay = Path(tmp) / "m.csv", Path(tmp) / "k.csv", Path(tmp) / "a.csv"
            m.write_text("Stok Kodu;Ürün;Stok;Fiyat\nA;a;5;100,00\nA;a;6;100\nB;b;yok;10\nC;c;-2;pahalı\n", encoding="utf-8")
            k.write_text("Kanal;Barkod;Stok;Fiyat\nK;A;5;100\nK;a;5;100\nK;C;var;x\n;A;1;1\n", encoding="utf-8")
            ay.write_text("Kanal;Fiyat Farkı %\nK;5\nL;çok\n", encoding="utf-8")
            mr, kl, a, uy = main.oku(m, k, ay)
            self.assertEqual(({s: (x["stok"], x["fiyat"]) for s, x in mr.items()}, [(x["sku"], x["stok"], x["fiyat"]) for x in kl], a),
                             ({"a": (D(5), D("100.00")), "c": (D(0), None)}, [("A", D(5), D(100)), ("C", None, None)], {"k": D(5)}))
            self.assertEqual(sorted(u["tur"] for u in uy), ["Kanal değeri okunamadı", "Kanal değeri okunamadı", "Kanal fiyat farkı okunamadı", "Merkez fiyatı okunamadı", "Merkez stoku okunamadı",
                                                            "Mükerrer kanal kaydı", "Mükerrer stok kodu"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--merkez", str(m), "--kanal", str(k), "--cikti", c]
            self.assertEqual(main.main(arg + ["--kanal-ayar", str(ay), "--fiyat-tolerans", "2,5", "--stok-tolerans", "1", "--tam-stok"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(arg + ["--fiyat-tolerans", "-1"]), 2)
            self.assertEqual(main.main(arg + ["--kanal-ayar", str(Path(tmp) / "yok.csv")]), 1)
            k.write_text("Kanal;Stok Kodu\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
