"""
Parsiyel Yük Konsolidasyon Planı — Workers / Workless kod bloğu
Lojistik ve Taşımacılık › Parsiyel Operasyon › Parsiyel Operasyon Uzmanı

Bekleyen parsiyel yükleri varış bölgesine göre gruplar ve araçlara yerleştirir; her aracın ağırlık, hacim ve
yükleme metresi (LDM) doluluğunu hesaplar:
  - LDM = ⌈palet adedi ÷ istif katı⌉ × palet eni × palet boyu ÷ 2,4 (standart dorse iç genişliği 2,4 m)
  - Ödenebilir ağırlık = en büyüğü (brüt kg, m³ × 333, LDM × 1.750) — oranlar değiştirilebilir
  - Yerleştirme: bölgeler en erken son yükleme tarihine göre sırayla; bölge içinde son yükleme tarihi, sonra büyük
    yük önce; açık araçlardan sığdığı ilkine (first-fit), sığmazsa kalan en büyük uygun araç açılır. Bölge bitince
    her araç, yükünü taşıyabilen en küçük müsait araç tipine indirilir.
  - ADR'li yük yalnız ADR donanımlı araca konur.
  - Kontroller: hazır olmayan yük, son yükleme tarihi geçmiş yük, hiçbir araca sığmayan yük, araç yetmediği için
    bekleyen yük, düşük doluluklu araç (varsayılan %60 altı).
Rapor: araç planı, yükleme listesi, bekleyen yükler, bölge özeti, uyarılar. İnternete bağlanmaz.

Kullanım:
    python main.py                                                     # örnek: 25 yük, 4 bölge, plan tarihi 09.10.2026
    python main.py --yukler yukler.xlsx --araclar araclar.csv --tarih 09.10.2026 --min-doluluk 70
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
DORSE_GENISLIK = 2.4
EVET = ("evet", "e", "x", "var", "yes", "1")

YUK_SUTUNLARI = {"no": ("yuk no", "no", "siparis no", "referans"), "musteri": ("musteri", "gonderici"),
                 "bolge": ("varis bolgesi", "bolge", "hat"), "sehir": ("varis sehri", "sehir", "varis"),
                 "palet": ("palet adedi", "palet", "kap adedi"), "en": ("palet en m", "palet en", "en m"), "boy": ("palet boy m", "palet boy", "boy m"),
                 "istif": ("istif kati", "istif", "istiflenebilir kat"), "kg": ("brut agirlik kg", "brut agirlik", "agirlik kg", "kg"),
                 "m3": ("hacim m3", "hacim m", "hacim", "m3"), "hazir": ("hazir tarihi", "hazir", "depoya giris"),
                 "son": ("son yukleme tarihi", "son yukleme", "termin"), "adr": ("adr", "tehlikeli madde"), "navlun": ("navlun eur", "navlun", "ucret")}
ARAC_SUTUNLARI = {"tip": ("arac tipi", "tip", "arac"), "adet": ("adet", "musait adet"), "kg": ("yuk kapasitesi kg", "yuk kapasitesi", "kapasite kg"),
                  "m3": ("hacim m3", "hacim m", "hacim"), "ldm": ("ldm", "yukleme metresi"), "adr": ("adr",)}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def sayi(x, varsayilan=None) -> float | None:
    if x in (None, ""):
        return varsayilan
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip().replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return varsayilan


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def binlik(x: float, ondalik: int = 0) -> str:
    return f"{x:,.{ondalik}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Yuk:
    satir: int
    no: str
    musteri: str
    bolge: str
    sehir: str
    palet: int
    en: float
    boy: float
    istif: int
    kg: float
    m3: float
    hazir: date | None
    son: date | None
    adr: bool
    navlun: float | None
    arac: str = ""
    durum: str = ""
    odenebilir: float = 0.0

    @property
    def ldm(self) -> float:
        return math.ceil(self.palet / max(self.istif, 1)) * self.en * self.boy / DORSE_GENISLIK if self.palet else 0.0


@dataclass
class AracTipi:
    tip: str
    adet: int
    kg: float
    m3: float
    ldm: float
    adr: bool

    def sigar(self, kg: float, m3: float, ldm: float, adr: bool) -> bool:
        return kg <= self.kg + 1e-9 and m3 <= self.m3 + 1e-9 and ldm <= self.ldm + 1e-9 and (self.adr or not adr)


@dataclass
class Arac:
    no: str
    tip: AracTipi
    bolge: str
    yukler: list = field(default_factory=list)

    @property
    def kg(self) -> float:
        return sum(y.kg for y in self.yukler)

    @property
    def m3(self) -> float:
        return sum(y.m3 for y in self.yukler)

    @property
    def ldm(self) -> float:
        return sum(y.ldm for y in self.yukler)

    @property
    def adr(self) -> bool:
        return any(y.adr for y in self.yukler)

    def alir(self, y: Yuk) -> bool:
        return self.tip.sigar(self.kg + y.kg, self.m3 + y.m3, self.ldm + y.ldm, self.adr or y.adr)

    def oranlar(self) -> dict[str, float]:
        return {"Ağırlık": self.kg / self.tip.kg, "Hacim": self.m3 / self.tip.m3, "LDM": self.ldm / self.tip.ldm}

    @property
    def doluluk(self) -> float:
        return max(self.oranlar().values())

    @property
    def belirleyici(self) -> str:
        o = self.oranlar()
        return max(o, key=o.get)


def yukleri_oku(yol: Path) -> tuple[list[Yuk], list[str]]:
    sonuc, hatalar = [], []
    for r in kayitlar(yol, YUK_SUTUNLARI, ("no", "bolge", "kg")):
        if not r.get("no"):
            continue
        kg, m3 = sayi(r.get("kg")), sayi(r.get("m3"), 0.0)
        if kg is None:
            hatalar.append(f"Satır {r['_satir']} ({r['no']}): ağırlık okunamadı")
            continue
        sonuc.append(Yuk(r["_satir"], metin(r["no"]), metin(r.get("musteri")), metin(r.get("bolge")) or "—", metin(r.get("sehir")),
                         int(sayi(r.get("palet"), 0)), sayi(r.get("en"), 0.8), sayi(r.get("boy"), 1.2), int(sayi(r.get("istif"), 1) or 1), kg, m3,
                         tarih(r.get("hazir")), tarih(r.get("son")), katla(r.get("adr")) in EVET, sayi(r.get("navlun"))))
    return sonuc, hatalar


def araclari_oku(yol: Path) -> list[AracTipi]:
    sonuc = []
    for r in kayitlar(yol, ARAC_SUTUNLARI, ("tip", "kg")):
        if r.get("tip"):
            sonuc.append(AracTipi(metin(r["tip"]), int(sayi(r.get("adet"), 1)), sayi(r.get("kg"), 0), sayi(r.get("m3"), 1e9) or 1e9,
                                  sayi(r.get("ldm"), 13.6) or 13.6, katla(r.get("adr")) in EVET))
    if not sonuc:
        raise ValueError(f"{yol.name}: araç tipi yok")
    return sonuc


# ----------------------------------------------------------------------------
# Planlama
# ----------------------------------------------------------------------------

SIRA = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}


def planla(yukler: list[Yuk], tipler: list[AracTipi], plan_tarihi: date, min_doluluk: float = 0.6, hacim_kg: float = 333, ldm_kg: float = 1750,
           bekletme_gun: int = 2) -> dict:
    uyarilar = []

    def uyar(onem, tur, aciklama, yuk="", arac=""):
        uyarilar.append({"onem": onem, "tur": tur, "yuk": yuk, "arac": arac, "aciklama": aciklama})

    kalan = Counter({t.tip: t.adet for t in tipler})
    # ADR donanımlı araçlar ADR'siz yükler için en sona bırakılır
    buyukten = sorted(tipler, key=lambda t: (t.adr, -t.ldm, -t.kg, -t.m3))
    kucukten = sorted(tipler, key=lambda t: (t.adr, t.ldm, t.kg, t.m3))
    for y in yukler:
        y.odenebilir = max(y.kg, y.m3 * hacim_kg, y.ldm * ldm_kg)
    hazir = []
    for y in yukler:
        if y.hazir and y.hazir > plan_tarihi:
            y.durum = "Hazır değil"
            uyar("Bilgi", "Hazır değil", f"{y.no}: hazır tarihi {y.hazir:%d.%m.%Y}; bu plana alınmadı", y.no)
        elif not any(t.sigar(y.kg, y.m3, y.ldm, y.adr) for t in tipler):
            y.durum = "Sığmıyor"
            neden = []
            if y.kg > max(t.kg for t in tipler):
                neden.append(f"{binlik(y.kg)} kg")
            if y.m3 > max(t.m3 for t in tipler):
                neden.append(f"{binlik(y.m3, 1)} m³")
            if y.ldm > max(t.ldm for t in tipler):
                neden.append(f"{binlik(y.ldm, 1)} LDM")
            if y.adr and not any(t.adr for t in tipler):
                neden.append("ADR donanımlı araç yok")
            uyar("Yüksek", "Hiçbir araca sığmıyor", f"{y.no}: {', '.join(neden) or 'kapasite'}; komple veya özel araç gerekir", y.no)
        else:
            hazir.append(y)
            if y.son and y.son < plan_tarihi:
                uyar("Yüksek", "Son yükleme tarihi geçti", f"{y.no}: son yükleme {y.son:%d.%m.%Y}, plan tarihi {plan_tarihi:%d.%m.%Y}; öncelikli yüklenir, "
                     "müşteriyi bilgilendirin", y.no)

    bolgeler = defaultdict(list)
    for y in hazir:
        bolgeler[y.bolge].append(y)
    sira = sorted(bolgeler, key=lambda b: (min(y.son or date.max for y in bolgeler[b]), -sum(y.odenebilir for y in bolgeler[b]), b))
    araclar, sayac = [], 0
    ref_kg, ref_m3, ref_ldm = (max(getattr(t, a) for t in tipler) or 1 for a in ("kg", "m3", "ldm"))
    for b in sira:
        lst = sorted(bolgeler[b], key=lambda y: (not y.adr, y.son or date.max, -max(y.kg / ref_kg, y.m3 / ref_m3, y.ldm / ref_ldm)))
        bolge_araclari = []
        for y in lst:
            hedef = next((a for a in bolge_araclari if a.alir(y)), None)
            if hedef is None:
                tip = next((t for t in buyukten if kalan[t.tip] > 0 and t.sigar(y.kg, y.m3, y.ldm, y.adr)), None)
                if tip is None:
                    y.durum = "Araç yok"
                    continue
                kalan[tip.tip] -= 1
                sayac += 1
                hedef = Arac(f"A{sayac:02d}", tip, b)
                bolge_araclari.append(hedef)
            hedef.yukler.append(y)
            y.arac, y.durum = hedef.no, "Planlandı"
        for a in bolge_araclari:              # küçült
            for t in kucukten:
                if t is a.tip:
                    break
                if kalan[t.tip] > 0 and t.sigar(a.kg, a.m3, a.ldm, a.adr):
                    kalan[a.tip.tip] += 1
                    kalan[t.tip] -= 1
                    a.tip = t
                    break
        araclar.extend(bolge_araclari)
        bekleyen = [y for y in lst if y.durum == "Araç yok"]
        if bekleyen:
            uyar("Yüksek", "Araç yetmedi", f"{b}: {len(bekleyen)} yük için müsait araç kalmadı ({', '.join(y.no for y in bekleyen)}); ek araç kiralayın veya "
                 "sonraki sefere bırakın")
    for a in araclar:
        a.yukler.sort(key=lambda y: (y.sehir, y.no))
        if a.doluluk < min_doluluk:
            bekletilebilir = all(y.son and y.son >= plan_tarihi + timedelta(days=bekletme_gun) for y in a.yukler)
            uyar("Orta", "Düşük doluluk", f"{a.no} ({a.tip.tip}, {a.bolge}): en yüksek doluluk %{a.doluluk * 100:.0f} ({a.belirleyici}). "
                 + ("Tüm yüklerin son yükleme tarihi en az " + str(bekletme_gun) + " gün sonra; bekletip yeni yüklerle birleştirmeyi düşünün."
                    if bekletilebilir else "Son yükleme tarihi yakın yük var; komşu bölgeyle birleştirme veya küçük araç değerlendirin."), arac=a.no)
    uyarilar.sort(key=lambda u: (SIRA[u["onem"]], u["arac"], u["yuk"]))
    return {"araclar": araclar, "uyarilar": uyarilar, "kalan": kalan, "bolge_sirasi": sira}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KONTROL = PatternFill("solid", fgColor="FFF4CE")
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, basliklar, genislik=()):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def _doluluk_rengi(c, oran):
    c.number_format = "0%"
    c.fill = PatternFill("solid", fgColor="E3F4E1" if oran >= 0.85 else "FFF4CE" if oran >= 0.6 else "FDE2E1")


def rapor_yaz(cikti: Path, yukler: list[Yuk], tipler: list[AracTipi], s: dict, plan_tarihi: date) -> None:
    wb = Workbook()
    ap = wb.active
    ap.title = "Araç Planı"
    _baslik(ap, ["Araç", "Tip", "Bölge", "Şehirler", "Yük", "Palet", "Ağırlık (kg)", "Hacim (m³)", "LDM", "Ağırlık %", "Hacim %", "LDM %", "Doluluk",
                 "Belirleyici", "ADR", "Ödenebilir (kg)", "Navlun (EUR)", "Plaka / Not"], (6, 12, 14, 26, 5, 6, 11, 9, 7, 9, 8, 7, 8, 10, 5, 12, 11, 20))
    for a in s["araclar"]:
        o = a.oranlar()
        navlun = [y.navlun for y in a.yukler if y.navlun is not None]
        ap.append([a.no, a.tip.tip, a.bolge, ", ".join(dict.fromkeys(y.sehir for y in a.yukler)), len(a.yukler), sum(y.palet for y in a.yukler),
                   round(a.kg), round(a.m3, 1), round(a.ldm, 2), o["Ağırlık"], o["Hacim"], o["LDM"], a.doluluk, a.belirleyici, "Evet" if a.adr else "",
                   round(sum(y.odenebilir for y in a.yukler)), sum(navlun) if navlun else None, ""])
        r = ap.max_row
        for j in (10, 11, 12):
            ap.cell(r, j).number_format = "0%"
        _doluluk_rengi(ap.cell(r, 13), a.doluluk)
        ap.cell(r, 18).fill = KONTROL
    ap.append([])
    ap.append([f"Plan tarihi {plan_tarihi:%d.%m.%Y} · kullanılmayan araçlar: " + (", ".join(f"{t} × {n}" for t, n in s["kalan"].items() if n > 0) or "yok")])

    yl = wb.create_sheet("Yükleme Listesi")
    _baslik(yl, ["Araç", "Yük No", "Müşteri", "Bölge", "Şehir", "Palet", "İstif", "Ağırlık (kg)", "Hacim (m³)", "LDM", "Ödenebilir (kg)", "Son Yükleme",
                 "ADR", "Navlun (EUR)"], (6, 8, 16, 14, 11, 6, 5, 10, 9, 6, 11, 11, 5, 10))
    for a in s["araclar"]:
        for y in a.yukler:
            yl.append([a.no, y.no, y.musteri, y.bolge, y.sehir, y.palet, y.istif, y.kg, y.m3, round(y.ldm, 2), round(y.odenebilir), y.son, "Evet" if y.adr else "",
                       y.navlun])
            yl.cell(yl.max_row, 12).number_format = "DD.MM.YYYY"
    yl.auto_filter.ref = f"A1:N{yl.max_row}"

    bk = wb.create_sheet("Bekleyen Yükler")
    _baslik(bk, ["Yük No", "Müşteri", "Bölge", "Şehir", "Ağırlık (kg)", "Hacim (m³)", "LDM", "Hazır", "Son Yükleme", "Neden", "Karar"],
            (8, 16, 14, 11, 10, 9, 6, 11, 11, 14, 24))
    for y in yukler:
        if y.durum != "Planlandı":
            bk.append([y.no, y.musteri, y.bolge, y.sehir, y.kg, y.m3, round(y.ldm, 2), y.hazir, y.son, y.durum, ""])
            bk.cell(bk.max_row, 8).number_format = bk.cell(bk.max_row, 9).number_format = "DD.MM.YYYY"
            bk.cell(bk.max_row, 11).fill = KONTROL

    bo = wb.create_sheet("Bölge Özeti")
    _baslik(bo, ["Bölge", "Yük", "Planlanan", "Bekleyen", "Ağırlık (kg)", "Hacim (m³)", "LDM", "Ödenebilir (kg)", "Araç", "Ort. Doluluk", "En Erken Son Yükleme"],
            (14, 6, 9, 9, 11, 9, 7, 12, 6, 10, 13))
    gruplar = defaultdict(list)
    for y in yukler:
        gruplar[y.bolge].append(y)
    for b, lst in sorted(gruplar.items()):
        ar = [a for a in s["araclar"] if a.bolge == b]
        bo.append([b, len(lst), sum(y.durum == "Planlandı" for y in lst), sum(y.durum != "Planlandı" for y in lst), round(sum(y.kg for y in lst)),
                   round(sum(y.m3 for y in lst), 1), round(sum(y.ldm for y in lst), 2), round(sum(y.odenebilir for y in lst)), len(ar),
                   sum(a.doluluk for a in ar) / len(ar) if ar else None, min((y.son for y in lst if y.son), default=None)])
        bo.cell(bo.max_row, 10).number_format = "0%"
        bo.cell(bo.max_row, 11).number_format = "DD.MM.YYYY"
    bo.append([])
    bo.append(["Araç tipleri"])
    for t in tipler:
        bo.append([t.tip, f"{t.adet} adet", f"{binlik(t.kg)} kg", f"{binlik(t.m3)} m³", f"{binlik(t.ldm, 1)} LDM", "ADR" if t.adr else ""])

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Yük", "Araç", "Açıklama", "Karar"], (9, 22, 8, 6, 100, 24))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["yuk"], u["arac"], u["aciklama"], ""])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 5).alignment = UST
        uy.cell(uy.max_row, 6).fill = KONTROL
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(yuk_yolu: Path, arac_yolu: Path, cikti: Path, plan_tarihi: date | None = None, min_doluluk: float = 0.6, hacim_kg: float = 333,
             ldm_kg: float = 1750) -> dict:
    plan_tarihi = plan_tarihi or date.today()
    yukler, hatalar = yukleri_oku(yuk_yolu)
    tipler = araclari_oku(arac_yolu)
    s = planla(yukler, tipler, plan_tarihi, min_doluluk, hacim_kg, ldm_kg)
    for h in hatalar:
        s["uyarilar"].append({"onem": "Orta", "tur": "Okunamayan satır", "yuk": "", "arac": "", "aciklama": h})
    rapor_yaz(cikti, yukler, tipler, s, plan_tarihi)
    return {**s, "yukler": yukler, "tipler": tipler, "plan_tarihi": plan_tarihi}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Parsiyel yükleri varış bölgesi, ağırlık, hacim ve LDM'ye göre araçlara gruplar; doluluk oranını hesaplar.")
    p.add_argument("--yukler", type=Path, default=ORNEK / "yukler.csv", help="Yük listesi (.xlsx/.csv)")
    p.add_argument("--araclar", type=Path, default=ORNEK / "araclar.csv", help="Araç Tipi, Adet, Yük Kapasitesi (kg), Hacim (m³), LDM, ADR")
    p.add_argument("--tarih", help="Plan (yükleme) tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 09.10.2026)")
    p.add_argument("--min-doluluk", type=float, default=60, help="Bu doluluk %%'sinin altındaki araç uyarılır (varsayılan 60)")
    p.add_argument("--hacim-kg", type=float, default=333, help="Ödenebilir ağırlıkta 1 m³ kaç kg sayılır (varsayılan 333)")
    p.add_argument("--ldm-kg", type=float, default=1750, help="Ödenebilir ağırlıkta 1 LDM kaç kg sayılır (varsayılan 1.750)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "konsolidasyon_plani.xlsx")
    a = p.parse_args(argv)
    for y in (a.yukler, a.araclar):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    plan = tarih(a.tarih) if a.tarih else (date(2026, 10, 9) if a.yukler == ORNEK / "yukler.csv" else None)
    if a.tarih and plan is None:
        print("[X] --tarih GG.AA.YYYY biçiminde olmalı")
        return 1
    try:
        s = calistir(a.yukler, a.araclar, a.cikti, plan, a.min_doluluk / 100, a.hacim_kg, a.ldm_kg)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    planli = sum(y.durum == "Planlandı" for y in s["yukler"])
    print(f"[OK] {len(s['yukler'])} yük · {planli} planlandı · {len(s['araclar'])} araç · plan tarihi {s['plan_tarihi']:%d.%m.%Y}")
    for x in s["araclar"]:
        print(f"[{'OK' if x.doluluk >= a.min_doluluk / 100 else '!'}] {x.no} {x.tip.tip:12} {x.bolge:14} {len(x.yukler):2} yük · doluluk %{x.doluluk * 100:.0f} "
              f"({x.belirleyici})")
    for u in s["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
