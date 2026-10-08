import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
PLAN = date(2026, 10, 9)


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "p.xlsx"
        cls.s = main.calistir(ORNEK / "yukler.csv", ORNEK / "araclar.csv", cls.cikti, PLAN)
        cls.y = {y.no: y for y in cls.s["yukler"]}
        cls.u = {}
        for u in cls.s["uyarilar"]:
            cls.u.setdefault(u["tur"], []).append(u)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_ldm_ve_odenebilir(self):
        y = self.y["Y-001"]                                   # 6 palet, istif 2 → 3 taban × 0,96 m² ÷ 2,4
        self.assertAlmostEqual(y.ldm, 1.2)
        self.assertAlmostEqual(y.odenebilir, 10.2 * 333)
        self.assertAlmostEqual(self.y["Y-025"].ldm, 12)

    def test_plan(self):
        self.assertEqual([(a.no, a.tip.tip, a.bolge, len(a.yukler)) for a in self.s["araclar"]],
                         [("A01", "Tenteli TIR", "Almanya-Batı", 6), ("A02", "Kamyon", "Almanya-Batı", 2), ("A03", "Mega TIR", "Almanya-Güney", 8),
                          ("A04", "Kamyon", "Benelüks", 3), ("A05", "Tenteli TIR", "Fransa", 4)])
        self.assertEqual(self.y["Y-004"].arac, "A03")          # ADR yük ADR'li araçta
        for a in self.s["araclar"]:
            self.assertLessEqual(a.kg, a.tip.kg)
            self.assertLessEqual(a.m3, a.tip.m3)
            self.assertLessEqual(a.ldm, a.tip.ldm + 1e-9)
        self.assertEqual({y.durum for y in self.s["yukler"] if not y.arac}, {"Hazır değil", "Sığmıyor"})

    def test_uyarilar(self):
        self.assertIn("26.000 kg", self.u["Hiçbir araca sığmıyor"][0]["aciklama"])
        self.assertEqual(self.u["Son yükleme tarihi geçti"][0]["yuk"], "Y-012")
        self.assertEqual(self.u["Hazır değil"][0]["yuk"], "Y-006")
        self.assertEqual([u["arac"] for u in self.u["Düşük doluluk"]], ["A05"])
        self.assertIn("bekletip", self.u["Düşük doluluk"][0]["aciklama"])

    def test_excel(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Araç Planı", "Yükleme Listesi", "Bekleyen Yükler", "Bölge Özeti", "Uyarılar"])
        self.assertEqual(wb["Araç Planı"]["N4"].value, "Ağırlık")
        self.assertEqual(wb["Bekleyen Yükler"].max_row, 3)


class KuralTesti(unittest.TestCase):
    def test_kucultme_ve_arac_yetmedi(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t / "a.csv").write_text("Araç Tipi;Adet;Yük Kapasitesi (kg);Hacim (m³);LDM;ADR\nTIR;1;24000;90;13,6;Hayır\nKamyon;1;10000;45;7,2;Hayır\n",
                                     encoding="utf-8")
            (t / "y.csv").write_text("Yük No;Varış Bölgesi;Palet Adedi;Brüt Ağırlık (kg);Hacim (m³);Son Yükleme Tarihi;ADR\n"
                                     "A;X;4;3000;6;10.10.2026;Hayır\nB;Y;20;20000;40;12.10.2026;Hayır\n"
                                     "C;Y;10;9000;20;12.10.2026;Hayır\nD;Z;2;1000;3;;Evet\n", encoding="utf-8")
            s = main.calistir(t / "y.csv", t / "a.csv", t / "o.xlsx", PLAN)
            y = {x.no: x for x in s["yukler"]}
            self.assertEqual(y["D"].durum, "Sığmıyor")                       # ADR araç yok
            tipler = {a.bolge: a.tip.tip for a in s["araclar"]}
            self.assertEqual(tipler["X"], "Kamyon")                          # TIR açılıp küçültüldü
            self.assertEqual(y["C"].durum, "Araç yok")
            self.assertIn("Araç yetmedi", {u["tur"] for u in s["uyarilar"]})

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--araclar", str(Path(tmp) / "yok.csv")]), 1)
            self.assertEqual(main.main(["--tarih", "yarın", "--cikti", str(Path(tmp) / "x.xlsx")]), 1)


if __name__ == "__main__":
    unittest.main()
