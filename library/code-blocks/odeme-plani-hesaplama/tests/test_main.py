import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def satirlar(p):
    return [(s["taksit"], s["anapara"], s["faiz"], s["kkdf"], s["bsmv"], s["kalan"]) for s in p["satirlar"]]


class PlanTesti(unittest.TestCase):
    def test_esit_taksit(self):
        p = main.plan(D(1000), D(10), 2)
        self.assertEqual(satirlar(p), [(D("576.19"), D("476.19"), D("100.00"), D("0.00"), D("0.00"), D("523.81")), (D("576.19"), D("523.81"), D("52.38"), D("0.00"), D("0.00"), D("0.00"))])
        self.assertEqual((p["toplam"], p["ilk_taksit"], p["son_taksit"], p["balon"]), ({"taksit": D("1152.38"), "anapara": D("1000.00"), "faiz": D("152.38"), "kkdf": D("0.00"), "bsmv": D("0.00")}, D("576.19"), D("576.19"), None))
        p = main.plan(D(1000), D(10), 2, bsmv=D(10), kkdf=D(10))
        self.assertEqual(satirlar(p), [(D("591.70"), D("471.70"), D("100.00"), D("10.00"), D("10.00"), D("528.30")), (D("591.69"), D("528.30"), D("52.83"), D("5.28"), D("5.28"), D("0.00"))])
        self.assertEqual(p["toplam"], {"taksit": D("1183.39"), "anapara": D("1000.00"), "faiz": D("152.83"), "kkdf": D("15.28"), "bsmv": D("15.28")})
        p = main.plan(D(1200), D(0), 12, bsmv=D(15), kkdf=D(15))
        self.assertEqual(({s["taksit"] for s in p["satirlar"]}, p["toplam"]["faiz"], p["satirlar"][-1]["kalan"]), ({D(100)}, D(0), D(0)))
        p = main.plan(D(1000), D(0), 3)
        self.assertEqual([s["taksit"] for s in p["satirlar"]], [D("333.33"), D("333.33"), D("333.34")])
        p = main.plan(D(100000), D("3.5"), 12, bsmv=D(15), kkdf=D(15), ilk=date(2026, 1, 31))
        beklenen = round(100000 * 0.0455 / (1 - 1.0455 ** -12), 2)
        s = p["satirlar"]
        self.assertAlmostEqual(float(p["ilk_taksit"]), beklenen, places=2)
        self.assertEqual((s[0]["faiz"], s[0]["kkdf"], s[0]["bsmv"], s[0]["anapara"], len(s), s[-1]["kalan"], p["toplam"]["anapara"]), (D("3500.00"), D("525.00"), D("525.00"), p["ilk_taksit"] - D(4550), 12, D(0), D(100000)))
        self.assertTrue(all(a["taksit"] == p["ilk_taksit"] for a in s[:-1]) and abs(s[-1]["taksit"] - p["ilk_taksit"]) < D("0.20") and all(a["anapara"] < b["anapara"] for a, b in zip(s, s[1:])))
        self.assertEqual([x["tarih"] for x in s][:4] + [s[-1]["tarih"]], [date(2026, 1, 31), date(2026, 2, 28), date(2026, 3, 31), date(2026, 4, 30), date(2026, 12, 31)])

    def test_esit_anapara_ve_balon(self):
        p = main.plan(D(1000), D(10), 4, "esit-anapara")
        self.assertEqual(([(s["taksit"], s["anapara"], s["faiz"], s["kalan"]) for s in p["satirlar"]], p["toplam"]["faiz"]),
                         ([(D("350.00"), D("250.00"), D("100.00"), D("750.00")), (D("325.00"), D("250.00"), D("75.00"), D("500.00")), (D("300.00"), D("250.00"), D("50.00"), D("250.00")), (D("275.00"), D("250.00"), D("25.00"), D("0.00"))], D("250.00")))
        p = main.plan(D(500000), D(3), 24, "esit-anapara", bsmv=D(5))
        s = p["satirlar"]
        self.assertEqual((s[0]["taksit"], s[0]["anapara"], s[0]["faiz"], s[0]["bsmv"], s[0]["kkdf"], s[-1]["anapara"], s[-1]["kalan"], p["toplam"]["anapara"]),
                         (D("36583.33"), D("20833.33"), D("15000.00"), D("750.00"), D("0.00"), D("20833.41"), D("0.00"), D("500000.00")))
        self.assertTrue(abs(p["toplam"]["faiz"] - D(187500)) < D("0.50") and all(a["taksit"] > b["taksit"] for a, b in zip(s[:-2], s[1:-1])))
        p = main.plan(D(1000), D(10), 2, "balon", D(400))
        self.assertEqual((satirlar(p), p["balon"]), ([(D("385.71"), D("285.71"), D("100.00"), D("0.00"), D("0.00"), D("714.29")), (D("785.72"), D("714.29"), D("71.43"), D("0.00"), D("0.00"), D("0.00"))], D(400)))
        p = main.plan(D(1000), D(0), 4, "balon", D(200))
        self.assertEqual([s["taksit"] for s in p["satirlar"]], [D(200), D(200), D(200), D(400)])
        p = main.plan(D(800000), D("2.8"), 36, "balon", D(300000), D(15), D(15))
        s = p["satirlar"]
        self.assertEqual((s[0]["faiz"], s[0]["kkdf"], s[0]["bsmv"], s[-1]["kalan"], p["toplam"]["anapara"]), (D("22400.00"), D("3360.00"), D("3360.00"), D("0.00"), D("800000.00")))
        self.assertTrue(abs(s[-1]["taksit"] - s[0]["taksit"] - D(300000)) < D("0.50") and s[-2]["kalan"] > D(300000))

    def test_dogrulama(self):
        gecersiz = [(D(0), D(1), 12, "esit-taksit", None), (D(100), D(-1), 12, "esit-taksit", None), (D(100), D(100), 12, "esit-taksit", None), (D(100), D(1), 0, "esit-taksit", None), (D(100), D(1), 601, "esit-taksit", None),
                    (D(100), D(1), 12, "serbest", None), (D(100), D(1), 12, "balon", None), (D(100), D(1), 12, "balon", D(100)), (D(100), D(1), 12, "balon", D(0))]
        for tutar, faiz, vade, tip, balon in gecersiz:
            with self.assertRaises(ValueError):
                main.plan(tutar, faiz, vade, tip, balon)
        for bsmv, kkdf in ((D(-1), D(0)), (D(0), D(101))):
            with self.assertRaises(ValueError):
                main.plan(D(100), D(1), 12, bsmv=bsmv, kkdf=kkdf)
        self.assertEqual(len(main.plan(D(100), D(1), 1)["satirlar"]), 1)
        self.assertEqual(main.plan(D(100), D(1), 12, "esit-taksit", D(999))["balon"], None)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "krediler.csv", cikti)
            a, b, c = r["krediler"]
            self.assertEqual([(k["ad"], k["tip"], k["vade"], len(k["plan"]), k["toplam"]["anapara"], k["plan"][-1]["kalan"], k["plan"][0]["tarih"], k["plan"][-1]["tarih"]) for k in r["krediler"]],
                             [("Tüketici Kredisi", "esit-taksit", 12, 12, D(100000), D(0), date(2026, 11, 15), date(2027, 10, 15)), ("Ticari Taksitli Kredi", "esit-anapara", 24, 24, D(500000), D(0), date(2026, 11, 1), date(2028, 10, 1)),
                              ("Taşıt Kredisi (balonlu)", "balon", 36, 36, D(800000), D(0), date(2026, 11, 20), date(2029, 10, 20))])
            self.assertAlmostEqual(float(a["ilk_taksit"]), round(100000 * 0.0455 / (1 - 1.0455 ** -12), 2), places=2)
            self.assertEqual((a["masraf"], a["maliyet"], a["geri_odeme"], b["ilk_taksit"], c["balon"]), (D(500), a["toplam"]["faiz"] + a["toplam"]["kkdf"] + a["toplam"]["bsmv"] + 500, a["toplam"]["taksit"] + 500, D("36583.33"), D(300000)))
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]], [("Bilgi", "Balonlu plan", "Taşıt Kredisi (balonlu)")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Plan 1", "Plan 2", "Plan 3", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "k.csv"
            y.write_text("Açıklama;Anapara;Faiz Oranı;Taksit Sayısı;Tip;Balon;BSMV;KKDF;İlk Taksit;Ücret\nA;1.250.000,50;2,75;24;EŞİT ANAPARALI;;5;;x;\n;100;1;3;;;;;;\nB;100;1;2,5;;;;;;\nC;100;1;3;Serbest;;;;;\nD;100;1;3;Balonlu;;;;;\n"
                         "E;100;1;3;;;150;;;\nF;100;1;3;;;;;;-5\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(k["ad"], k["tutar"], k["faiz"], k["vade"], k["tip"], k["bsmv"], k["kkdf"], k["ilk"], k["masraf"]) for k in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("A", D("1250000.50"), D("2.75"), 24, "esit-anapara", D(5), D(0), None, D(0)), ("Kredi (satır 3)", D(100), D(1), 3, "esit-taksit", D(0), D(0), None, D(0))],
                              [("Tarih okunamadı", "A")] + [("Kredi satırı okunamadı", n) for n in "BCDEF"]))
            self.assertEqual([u["tur"] for u in main.hesapla(l)["uyarilar"]], ["Vergi ve fon oranı girilmedi"])
            c = str(Path(tmp) / "x.xlsx")
            self.assertEqual(main.main(["--krediler", str(y), "--cikti", c]), 0)
            tek = ["--tutar", "100.000", "--faiz", "3,5", "--vade", "12", "--cikti", c]
            self.assertEqual(main.main(tek + ["--bsmv", "15", "--kkdf", "15", "--ilk-taksit", "15.11.2026", "--masraf", "500"]), 0)
            self.assertEqual(load_workbook(c).sheetnames, ["Özet", "Plan 1", "Uyarılar"])
            self.assertEqual(main.main(["--tutar", "800000", "--faiz", "2.8", "--vade", "36", "--tip", "balon", "--balon", "300000", "--cikti", c]), 0)
            self.assertEqual(main.main(["--tutar", "500000", "--faiz", "3", "--vade", "24", "--tip", "esit-anapara", "--cikti", c]), 0)
            self.assertEqual([main.main(tek + v) for v in (["--bsmv", "x"], ["--kkdf", "101"], ["--ilk-taksit", "x"], ["--masraf", "-1"], ["--tip", "balon"], ["--krediler", str(y)])]
                             + [main.main(["--tutar", "100", "--cikti", c]), main.main(["--tutar", "0", "--faiz", "1", "--vade", "3", "--cikti", c]), main.main(["--tutar", "5", "--faiz", "100", "--vade", "3", "--cikti", c])], [2] * 9)
            self.assertEqual(main.main(["--krediler", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Kredi;Tutar\nA;1\n", "Tutar;Faiz;Vade\n;;\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(["--krediler", str(y), "--cikti", c]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
