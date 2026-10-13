"""
Ödeme Planı Hesaplama — Workers / Workless kod bloğu
Bankacılık Operasyonları › Kredi Operasyon Uzmanı

Kredi tutarı, faiz, vade ve ödeme tipine (eşit taksit, eşit anapara, balonlu) göre BSMV ve KKDF dahil ödeme planı çıkarır.
  Dönem faizi       = kalan anapara × aylık faiz oranı
  KKDF, BSMV        = dönem faizi × KKDF oranı;  dönem faizi × BSMV oranı
  Vergili oran (rv) = aylık faiz oranı × (1 + KKDF oranı + BSMV oranı)
  Eşit taksit       = anapara × rv ÷ (1 − (1 + rv)^−vade);   her taksitte anapara = taksit − faiz − KKDF − BSMV
  Eşit anapara      = her dönem anapara ÷ vade kadar anapara + o dönemin faizi ve vergileri
  Balonlu           = (anapara − balon ÷ (1 + rv)^vade) × rv ÷ (1 − (1 + rv)^−vade) eşit taksit; son taksitte ayrıca balon tutarı
  Son taksit        = kalan anaparanın tamamı + faizi ve vergileri (kuruş farkı son taksitte kapanır)
Dönemler eşit ve aylık kabul edilir; gün sayısına göre faiz, ödemesiz dönem ve ara ödeme yoktur. BSMV ve KKDF oranlarını siz verirsiniz.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
import calendar

SUTUN = {"ad": ("kredi", "kredi adi", "aciklama", "musteri"), "tutar": ("tutar", "kredi tutari", "anapara"), "faiz": ("aylik faiz", "aylik faiz orani", "faiz orani", "faiz"), "vade": ("vade ay", "vade", "taksit sayisi"),
         "tip": ("odeme tipi", "tip", "plan tipi"), "balon": ("balon tutari", "balon", "balon odeme"), "bsmv": ("bsmv",), "kkdf": ("kkdf",), "ilk": ("ilk taksit tarihi", "ilk taksit", "ilk odeme tarihi"), "masraf": ("masraf", "tahsis ucreti", "ucret")}
TIP = {"esit taksit": "esit-taksit", "esit taksitli": "esit-taksit", "esit-taksit": "esit-taksit", "esit anapara": "esit-anapara", "esit anaparali": "esit-anapara", "esit-anapara": "esit-anapara", "balon": "balon", "balonlu": "balon",
       "balon odemeli": "balon", "": "esit-taksit"}
TIP_AD = {"esit-taksit": "Eşit taksit", "esit-anapara": "Eşit anapara", "balon": "Balonlu"}
Z, YUZ, BIR = Decimal(0), Decimal(100), Decimal(1)


def ay_ekle(t: date, n: int) -> date:
    y, a = divmod(t.year * 12 + t.month - 1 + n, 12)
    return date(y, a + 1, min(t.day, calendar.monthrange(y, a + 1)[1]))


def olcu(x) -> Decimal | None:
    """Oran okur: virgül ya da tek nokta ondalık ayırıcıdır (3,5 = 3.5)."""
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x or ""))
    try:
        return Decimal(s.replace(".", "").replace(",", ".") if "," in s else s) if s else None
    except InvalidOperation:
        return None


def dogrula(tutar, faiz, vade, tip, balon, bsmv, kkdf) -> str | None:
    if tutar is None or tutar <= 0:
        return "kredi tutarı pozitif olmalı"
    if faiz is None or not 0 <= faiz < 100:
        return "aylık faiz oranı 0 ile 100 arasında olmalı"
    if vade is None or not 1 <= vade <= 600:
        return "vade 1 ile 600 ay arasında tam sayı olmalı"
    if tip not in TIP_AD:
        return "ödeme tipi Eşit Taksit, Eşit Anapara ya da Balon olmalı"
    if any(o is None or not 0 <= o <= 100 for o in (bsmv, kkdf)):
        return "BSMV ve KKDF oranı 0 ile 100 arasında olmalı"
    if tip == "balon" and (balon is None or not 0 < balon < tutar):
        return "balonlu planda balon tutarı sıfırdan büyük ve kredi tutarından küçük olmalı"
    return None


def plan(tutar: Decimal, faiz: Decimal, vade: int, tip: str = "esit-taksit", balon: Decimal | None = None, bsmv: Decimal = Z, kkdf: Decimal = Z, ilk: date | None = None) -> dict:
    hata = dogrula(tutar, faiz, vade, tip, balon, bsmv, kkdf)
    if hata:
        raise ValueError(hata)
    r = faiz / YUZ
    rv = r * (BIR + (bsmv + kkdf) / YUZ)
    carpan = (BIR + rv) ** vade
    balon = balon if tip == "balon" else Z
    if tip == "esit-anapara":
        sabit_anapara, taksit = yuvarla(tutar / vade), None
    else:
        taksit = yuvarla((tutar - balon / carpan) * rv / (BIR - BIR / carpan) if rv else (tutar - balon) / vade)
    bakiye, satirlar = tutar, []
    for i in range(1, vade + 1):
        f = yuvarla(bakiye * r)
        k, b = yuvarla(f * kkdf / YUZ), yuvarla(f * bsmv / YUZ)
        anapara = bakiye if i == vade else sabit_anapara if tip == "esit-anapara" else taksit - f - k - b
        if anapara < 0 or anapara > bakiye:
            raise ValueError("bu faiz, vade ve balon tutarıyla taksit dönem faizini karşılamıyor; değerleri kontrol edin")
        bakiye -= anapara
        satirlar.append({"no": i, "tarih": ay_ekle(ilk, i - 1) if ilk else None, "taksit": anapara + f + k + b, "anapara": anapara, "faiz": f, "kkdf": k, "bsmv": b, "kalan": bakiye})
    t = {a: sum((s[a] for s in satirlar), Z) for a in ("taksit", "anapara", "faiz", "kkdf", "bsmv")}
    return {"satirlar": satirlar, "toplam": t, "ilk_taksit": satirlar[0]["taksit"], "son_taksit": satirlar[-1]["taksit"], "balon": balon or None}


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("tutar", "faiz", "vade")):
        if not metin(r["tutar"]) and not metin(r["ad"]):
            continue
        ad = metin(r["ad"]) or f"Kredi (satır {r['_satir']})"
        vade = olcu(r["vade"])
        k = {"ad": ad, "tutar": para(r["tutar"]), "faiz": olcu(r["faiz"]), "vade": int(vade) if vade is not None and vade == vade.to_integral_value() else None, "tip": TIP.get(katla(r["tip"])), "balon": para(r["balon"]),
             "bsmv": olcu(r["bsmv"]) if metin(r["bsmv"]) else Z, "kkdf": olcu(r["kkdf"]) if metin(r["kkdf"]) else Z, "ilk": tarih(r["ilk"]), "masraf": para(r["masraf"]) if metin(r["masraf"]) else Z}
        hata = dogrula(k["tutar"], k["faiz"], k["vade"], k["tip"], k["balon"], k["bsmv"], k["kkdf"]) or ("masraf sıfır veya pozitif olmalı" if k["masraf"] is None or k["masraf"] < 0 else None)
        if hata:
            uy.append(uyari("Yüksek", "Kredi satırı okunamadı", ad, f"Satır {r['_satir']}: {hata}. Satır alınmadı."))
            continue
        if metin(r["ilk"]) and k["ilk"] is None:
            uy.append(uyari("Orta", "Tarih okunamadı", ad, f"İlk taksit tarihi '{metin(r['ilk'])}' okunamadı; plan tarihsiz üretildi."))
        l.append(k)
    return l, uy


def hesapla(krediler: list[dict]) -> dict:
    if not krediler:
        raise ValueError("Hesaplanacak kredi yok.")
    uy = []
    for k in krediler:
        try:
            p = plan(k["tutar"], k["faiz"], k["vade"], k["tip"], k["balon"], k["bsmv"], k["kkdf"], k["ilk"])
        except ValueError as h:
            raise ValueError(f"{k['ad']}: {h}") from None
        k.update(plan=p["satirlar"], toplam=p["toplam"], ilk_taksit=p["ilk_taksit"], son_taksit=p["son_taksit"])
        k["maliyet"] = p["toplam"]["faiz"] + p["toplam"]["kkdf"] + p["toplam"]["bsmv"] + k["masraf"]
        k["geri_odeme"] = p["toplam"]["taksit"] + k["masraf"]
        if k["bsmv"] == 0 and k["kkdf"] == 0 and k["faiz"] > 0:
            uy.append(uyari("Bilgi", "Vergi ve fon oranı girilmedi", k["ad"], "BSMV ve KKDF oranı sıfır; plan yalnız anapara ve faiz içeriyor. Kredi türünüz için güncel oranları girin."))
        if k["tip"] == "balon":
            uy.append(uyari("Bilgi", "Balonlu plan", k["ad"], f"Son taksit {p['son_taksit']:,.2f}; bunun {k['balon']:,.2f} tutarı balon ödemedir."))
    return {"krediler": krediler, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    _baslik(ws, ["Kredi", "Ödeme Tipi", "Kredi Tutarı", "Aylık Faiz %", "Vade (ay)", "BSMV %", "KKDF %", "İlk Taksit", "Son Taksit", "Toplam Faiz", "Toplam KKDF", "Toplam BSMV", "Masraf", "Toplam Maliyet", "Toplam Geri Ödeme", "Plan Sayfası"],
            (30, 14, 16, 10, 9, 8, 8, 14, 14, 15, 14, 14, 12, 16, 17, 11))
    for i, k in enumerate(r["krediler"], 1):
        t = k["toplam"]
        _satir(ws, [k["ad"], TIP_AD[k["tip"]], k["tutar"], k["faiz"], k["vade"], k["bsmv"], k["kkdf"], k["ilk_taksit"], k["son_taksit"], t["faiz"], t["kkdf"], t["bsmv"], k["masraf"] or None, k["maliyet"], k["geri_odeme"], f"Plan {i}"],
               para_s=(3, 4, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15))
    ws.append([])
    ws.append(["Toplam maliyet = faiz + KKDF + BSMV + masraf. Dönemler eşit ve aylık kabul edilmiştir; gerçek planda gün sayısı, tatil kaydırması ve kuruş yuvarlaması nedeniyle küçük farklar olabilir."])
    for i, k in enumerate(r["krediler"], 1):
        ws = wb.create_sheet(f"Plan {i}")
        ws.append([f"{k['ad']} · {TIP_AD[k['tip']]} · {k['tutar']:,.2f} · aylık %{k['faiz']:f} · {k['vade']} ay · BSMV %{k['bsmv']:f} · KKDF %{k['kkdf']:f}"])
        _baslik(ws, ["Taksit No", "Tarih", "Taksit Tutarı", "Anapara", "Faiz", "KKDF", "BSMV", "Kalan Anapara"], (10, 12, 16, 16, 15, 13, 13, 17), dondur="A3")
        for s in k["plan"]:
            _satir(ws, [s["no"], s["tarih"], s["taksit"], s["anapara"], s["faiz"], s["kkdf"], s["bsmv"], s["kalan"]], para_s=(3, 4, 5, 6, 7, 8), tarih_s=(2,))
        t = k["toplam"]
        n = _satir(ws, ["TOPLAM", None, t["taksit"], t["anapara"], t["faiz"], t["kkdf"], t["bsmv"], None], para_s=(3, 4, 5, 6, 7))
        for h in ws[n]:
            h.font = Font(bold=True)
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path) -> dict:
    krediler, uy = oku(yol)
    r = hesapla(krediler)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Kredi tutarı, faiz, vade ve ödeme tipine göre BSMV ve KKDF dahil ödeme planı çıkarır.")
    p.add_argument("--krediler", type=Path, help="Toplu hesap: Kredi, Tutar, Aylık Faiz, Vade (ay), Ödeme Tipi, Balon Tutarı, BSMV, KKDF, İlk Taksit Tarihi, Masraf")
    p.add_argument("--tutar", help="Tek kredi: kredi tutarı")
    p.add_argument("--faiz", help="Tek kredi: aylık faiz oranı, yüzde (örn. 3,5)")
    p.add_argument("--vade", type=int, help="Tek kredi: vade, ay")
    p.add_argument("--tip", choices=tuple(TIP_AD), default="esit-taksit", help="Ödeme tipi (varsayılan esit-taksit)")
    p.add_argument("--balon", help="Balonlu planda son taksitle ödenecek balon tutarı")
    p.add_argument("--bsmv", default="0", help="BSMV oranı, yüzde (varsayılan 0; güncel oranı siz girin)")
    p.add_argument("--kkdf", default="0", help="KKDF oranı, yüzde (varsayılan 0; güncel oranı siz girin)")
    p.add_argument("--ilk-taksit", help="İlk taksit tarihi (GG.AA.YYYY); verilmezse plan tarihsiz olur")
    p.add_argument("--masraf", default="0", help="Tahsis ücreti gibi peşin masraf (toplam maliyete eklenir)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "odeme_plani.xlsx")
    a = p.parse_args(argv)
    tekli = any(x is not None for x in (a.tutar, a.faiz, a.vade))
    if tekli and (a.krediler or None in (a.tutar, a.faiz, a.vade)):
        print("[X] Tek kredi için --tutar, --faiz ve --vade birlikte verilmeli; --krediler ile aynı anda kullanılamaz")
        return 2
    try:
        if tekli:
            k = {"ad": "Kredi", "tutar": para(a.tutar), "faiz": olcu(a.faiz), "vade": a.vade, "tip": a.tip, "balon": para(a.balon), "bsmv": olcu(a.bsmv), "kkdf": olcu(a.kkdf), "ilk": tarih(a.ilk_taksit), "masraf": para(a.masraf)}
            hata = dogrula(k["tutar"], k["faiz"], k["vade"], k["tip"], k["balon"], k["bsmv"], k["kkdf"]) or ("--ilk-taksit GG.AA.YYYY biçiminde olmalı" if a.ilk_taksit and k["ilk"] is None else None) or (
                "--masraf sıfır veya pozitif olmalı" if k["masraf"] is None or k["masraf"] < 0 else None)
            if hata:
                print(f"[X] {hata}")
                return 2
            r = hesapla([k])
            rapor_yaz(a.cikti, r)
        else:
            yol = a.krediler or ORNEK / "krediler.csv"
            if not yol.exists():
                print(f"[X] Dosya bulunamadı: {yol}")
                return 1
            r = calistir(yol, a.cikti)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    for k in r["krediler"]:
        print(f"[OK] {k['ad']}: {TIP_AD[k['tip']]}, {k['vade']} ay · ilk taksit {k['ilk_taksit']:,.2f} · son taksit {k['son_taksit']:,.2f} · toplam geri ödeme {k['geri_odeme']:,.2f} · toplam maliyet {k['maliyet']:,.2f}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
