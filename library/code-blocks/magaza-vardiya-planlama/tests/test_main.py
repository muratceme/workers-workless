import sys
import tempfile
import unittest
from datetime import time
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def vd(ad, basla, bitis, mola=60):
    return {"ad": ad, "basla": D(basla), "bitis": D(bitis), "mola": D(mola), "net": D(bitis) - D(basla) - D(mola) / 60, "etiket": ad}


def ps(*limitler, izin=None):
    return [{"ad": f"P{i}", "limit": D(l), "izin": set((izin or {}).get(i, ()))} for i, l in enumerate(limitler)]


def tr(saatler, gunler=(0,), musteri=10):
    return {g: {h: D(musteri) for h in saatler} for g in gunler}


SABAH, AKSAM, ERKEN = vd("Sabah", 10, 18), vd("Akşam", 14, 22), vd("Erken", 8, 16)


def atananlar(r):
    return sorted((g, f"P{i}", v["ad"]) for (i, g), v in r["atama"].items())


class PlanTesti(unittest.TestCase):
    def test_yardimcilar(self):
        self.assertEqual([main.gun_coz(x) for x in ("Pazar", "PAZARTESİ", "çarşamba", "Sal", "Paz", "12.10.2026", "x", "")], [6, 0, 2, 1, None, 0, None, None])
        self.assertEqual([main.saat_coz(x) for x in ("10", "10:30", "18.30", "09:00:00", time(14, 15), "25", "on")], [D(10), D("10.5"), D("18.5"), D(9), D("14.25"), None, None])
        self.assertEqual(([main.asgari_mola(D(x)) for x in ("4", "4.5", "7.5", "8")], main.saat_yaz(D("18.5"))), ([15, 30, 30, 60], "18:30"))

    def test_ihtiyac_ve_temel_atama(self):
        r = main.planla({0: {10: D(45), 11: D(10), 12: D(0), 13: D(40)}}, ps(45, 45, 45, 45), [SABAH])
        self.assertEqual([(k["saat"], k["ihtiyac"], k["atanan"], k["fark"]) for k in r["kapsama"]], [(10, 3, 3, 0), (11, 2, 3, 1), (12, 2, 3, 1), (13, 2, 3, 1)])      # 45 ÷ 20 → 3; en az 2
        self.assertEqual((atananlar(r), r["toplam"]), ([(0, "P0", "Sabah"), (0, "P1", "Sabah"), (0, "P2", "Sabah")], {"vardiya": 3, "ihtiyac": 9, "atanan": 12, "eksik": 0, "fazla": 3, "saat": D(21)}))
        self.assertEqual(atananlar(main.planla(tr(range(10, 18)), ps(45, 45, 45, izin={0: {0}}), [SABAH])), [(0, "P1", "Sabah"), (0, "P2", "Sabah")])              # izinli personele atama yok
        r = main.planla(tr(range(10, 14)), ps(45, 45), [SABAH, vd("Kısa", 10, 14, 15)])
        self.assertEqual({a[2] for a in atananlar(r)}, {"Kısa"})                                                                                                   # aynı saatleri kapatan vardiyalardan kısa olan
        r = main.planla(tr(range(10, 22)), ps(45), [SABAH, AKSAM], min_personel=1)
        self.assertEqual((atananlar(r), r["ozet"][0]["eksik"], r["ozet"][0]["eksik_saat"]), ([(0, "P0", "Sabah")], 4, [18, 19, 20, 21]))                           # günde tek vardiya

    def test_kurallar(self):
        r = main.planla(tr(range(10, 18), (0, 1)), ps(7, 14), [SABAH], min_personel=1)
        self.assertEqual(atananlar(r), [(0, "P1", "Sabah"), (1, "P0", "Sabah")])                                    # kalan saati çok olan önce; haftalık sınır aşılmaz
        self.assertEqual(main.planla(tr(range(10, 18), (0, 1, 2)), ps(10), [SABAH], min_personel=1)["toplam"], {"vardiya": 1, "ihtiyac": 24, "atanan": 8, "eksik": 16, "fazla": 0, "saat": D(7)})
        self.assertEqual(len(main.planla(tr(range(10, 18), (0, 1, 2)), ps(45), [SABAH], min_personel=1, max_gun=2)["atama"]), 2)
        t = {0: {h: D(1) for h in range(18, 22)}, 1: {h: D(1) for h in range(8, 12)}}
        self.assertEqual(atananlar(main.planla(t, ps(45), [AKSAM, ERKEN], min_personel=1)), [(0, "P0", "Akşam")])    # 22:00 → 08:00 = 10 saat < 11
        self.assertEqual(len(main.planla(t, ps(45), [AKSAM, ERKEN], min_personel=1, dinlenme=D(10))["atama"]), 2)
        t[1] = {h: D(30) for h in range(8, 12)}                                                                       # salı daha yoğun: önce o planlanır, pazartesi akşamı sonraki güne bakılarak engellenir
        self.assertEqual(atananlar(main.planla(t, ps(45), [AKSAM, ERKEN], min_personel=1)), [(1, "P0", "Erken")])
        r = main.planla(tr(range(8, 20)), ps(45), [vd("Uzun", 8, 20), SABAH], min_personel=1, gunluk=D(9))
        self.assertEqual((atananlar(r), r["atlanan_vardiya"]), ([(0, "P0", "Sabah")], ["Uzun"]))
        for t2, p, v in (({}, ps(45), [SABAH]), (tr([10]), [], [SABAH]), (tr([10]), ps(45), [vd("Uzun", 8, 22)])):
            with self.assertRaises(ValueError):
                main.planla(t2, p, v)

    def test_plan_tablosu(self):
        r = main.planla(tr(range(10, 18), (0, 1)), ps(45, 45, izin={1: {1}}), [SABAH], min_personel=1)
        self.assertEqual([(p["personel"], p["gunler"], p["saat"], p["gun"], p["kalan"]) for p in r["plan"]], [("P0", {0: "", 1: "Sabah"}, D(7), 1, D(38)), ("P1", {0: "Sabah", 1: "İzin"}, D(7), 1, D(38))])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "v.xlsx"
            r = main.calistir(ORNEK / "saatlik_trafik.csv", ORNEK / "personel.csv", ORNEK / "vardiyalar.csv", cikti)
            self.assertEqual((r["toplam"], r["uyarilar"]), ({"vardiya": 30, "ihtiyac": 192, "atanan": 240, "eksik": 0, "fazla": 48, "saat": D(210)}, []))
            self.assertEqual([(o["gun"], o["vardiya"], o["ihtiyac"]) for o in r["ozet"]],
                             [("Pazartesi", 4, 24), ("Salı", 4, 24), ("Çarşamba", 4, 24), ("Perşembe", 4, 24), ("Cuma", 4, 24), ("Cumartesi", 5, 36), ("Pazar", 5, 36)])
            self.assertTrue(all(p["saat"] <= p["limit"] and p["gun"] <= 6 for p in r["plan"]))
            self.assertTrue(all(k["fark"] >= 0 for k in r["kapsama"]))
            self.assertEqual([(p["personel"], p["gunler"][0], p["gunler"][1]) for p in r["plan"] if "İzin" in p["gunler"].values()][5], ("Personel 6", "İzin", "İzin"))
            for (i, g), v in r["atama"].items():                                 # vardiyalar arası en az 11 saat
                sonra = r["atama"].get((i, g + 1))
                self.assertTrue(sonra is None or 24 - v["bitis"] + sonra["basla"] >= 11)
            self.assertEqual(load_workbook(cikti).sheetnames, ["Vardiya Planı", "Gün Özeti", "Saatlik Kapsama", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            t, p, v = Path(tmp) / "t.csv", Path(tmp) / "p.csv", Path(tmp) / "v.csv"
            t.write_text("Gün;Saat;Müşteri\nPazartesi;10:00;30\npazartesi;10;15\n13.10.2026;11:00;20\nBayram;10;5\nSalı;öğlen;5\n", encoding="utf-8")
            p.write_text("Personel;Haftalık Saat;İzin Günleri\nA;;Pazartesi, bayram\nB;60;\nC;az;13.10.2026\n", encoding="utf-8")
            v.write_text("Vardiya;Başlangıç;Bitiş;Mola\nSabah;10:00;18:00;30\nGece;22:00;06:00;60\nBozuk;sabah;18;\nKısa;10;14;\nHiç;10;11;60\n", encoding="utf-8")
            tf, pr, vr, uy = main.oku(t, p, v, D(45))
            self.assertEqual((tf, [(x["ad"], x["limit"], x["izin"]) for x in pr], [(x["ad"], x["net"], x["etiket"]) for x in vr]),
                             ({0: {10: D(45)}, 1: {11: D(20)}}, [("A", D(45), {0}), ("B", D(45), set()), ("C", D(45), {1})], [("Sabah", D("7.5"), "Sabah 10:00-18:00"), ("Kısa", D(4), "Kısa 10:00-14:00")]))
            self.assertEqual(sorted(u["tur"] for u in uy), ["Gece yarısını aşan vardiya", "Haftalık saat okunamadı", "Haftalık saat sınırın üzerinde", "Mola asgari sürenin altında", "Mola asgari sürenin altında",
                                                            "Trafik satırı okunamadı", "Trafik satırı okunamadı", "Vardiya okunamadı", "Vardiya okunamadı", "İzin günü tanınmadı"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--trafik", str(t), "--personel", str(p), "--vardiyalar", str(v), "--cikti", c]
            self.assertEqual(main.main(arg + ["--musteri-basina", "15", "--min-personel", "1", "--haftalik-saat", "40", "--gunluk-saat", "9", "--dinlenme-saat", "12", "--max-gun", "5"]), 0)
            r = main.calistir(t, p, v, Path(c), min_personel=4)
            self.assertIn("Karşılanamayan ihtiyaç", [u["tur"] for u in r["uyarilar"]])
            self.assertEqual(main.main(arg + ["--max-gun", "8"]), 2)
            self.assertEqual(main.main(arg + ["--musteri-basina", "0"]), 2)
            self.assertEqual(main.main(arg + ["--gunluk-saat", "3"]), 1)
            self.assertEqual(main.main(["--trafik", str(Path(tmp) / "yok.csv")]), 1)
            p.write_text("Personel\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
