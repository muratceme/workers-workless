import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "h.xlsx"
        cls.s = main.calistir(ORNEK / "sozlesme.csv", ORNEK / "faturalar.csv", cls.cikti, ORNEK / "kullanim.csv")
        cls.b = {(x.no, x.kalem): [b["tur"] for b in x.bulgular] for x in cls.s["faturalar"]}
        cls.m = {m["kalem"]: m for m in cls.s["mutabakat"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_mutabakat(self):
        o = self.m["Öğle yemeği"]
        self.assertEqual((o["hak"], o["faturalanan"], o["itiraz"]), (D(2789), D(2789), D(0)))   # asgari 120 garantisi uygulanmış
        self.assertGreater(o["hak"], o["fiili"])
        t = self.m["Temizlik personeli"]
        self.assertEqual((t["hak"], t["itiraz"]), (D("5.9000"), D("5800.00")))                    # 177 / 30; 0,1 × 58.000
        self.assertEqual(self.m["Güzergah 1 - Merkez"]["itiraz"], D("4800.00"))
        self.assertEqual(self.m["Fazla mesai"]["itiraz"], D("6240.00"))
        self.assertEqual(self.m["Güvenlik görevlisi"]["itiraz"], D("248000.00"))                  # mükerrer fatura

    def test_satir_bulgulari(self):
        self.assertEqual(self.b[("YEM2026000912", "Öğle yemeği")], [])
        self.assertEqual(self.b[("YEM2026000912", "Kumanya (gece vardiyası)")], ["Fiyat farkı"])
        self.assertIn("KDV oranı", self.b[("SRV2026000455", "Güzergah 2 - Gebze")])
        self.assertIn("Tevkifat tutarı", self.b[("TMZ2026000301", "Temizlik personeli")])
        self.assertIn("Hesap hatası", self.b[("GVN2026000188", "Fazla mesai")])
        self.assertEqual(self.b[("GVN2026000191", "Güvenlik görevlisi")], ["Mükerrer faturalama", "Fazla miktar"])
        self.assertEqual(self.b[("YEM2026000913", "Kahvaltı")], ["Sözleşmede yok"])

    def test_toplam_itiraz(self):
        # 1.000 hesap + 6.240 mesai + 248.000 mükerrer + 4.800 sefer + 5.800 puantaj + 1.176 fiyat + 9.000 sözleşmesiz
        self.assertEqual(sum((x.itiraz for x in self.s["faturalar"]), D(0)), D("276016.00"))
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["İtiraz Listesi", "Dönem Mutabakatı", "Fatura Kontrolü", "Kullanım Özeti", "Uyarılar"])
        it = wb["İtiraz Listesi"]
        self.assertEqual(it.cell(it.max_row, 8).value, 276016)


class KuralTesti(unittest.TestCase):
    def test_fiyat_donemi_ve_oran(self):
        f = main.sozlesme_oku(ORNEK / "sozlesme.csv")
        self.assertEqual(main.gecerli_fiyat(f, "Yemek", "öğle yemeği", date(2026, 6, 1)).fiyat, D(165))
        self.assertEqual(main.gecerli_fiyat(f, "yemek", "Öğle Yemeği", date(2026, 7, 1)).fiyat, D(185))
        self.assertIsNone(main.gecerli_fiyat(f, "Yemek", "Öğle yemeği", date(2027, 1, 1)))
        self.assertEqual(main.oran("9/10"), D("0.9"))
        self.assertEqual(main.oran("%20"), D("0.2"))
        self.assertEqual(main.donem_coz("Eylül 2026"), (2026, 9))
        self.assertEqual(main.donem_coz("09.2026"), (2026, 9))
        self.assertEqual(main.donem_coz("2026-09"), (2026, 9))

    def test_kullanimsiz_fatura_ve_faturasiz_kullanim(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t / "f.csv").write_text("Fatura No;Dönem;Hizmet;Kalem;Miktar;Birim Fiyat;Tutar\nA1;2026-09;Servis;Güzergah 1 - Merkez;10;2400;24000\n", encoding="utf-8")
            (t / "k.csv").write_text("Tarih;Hizmet;Kalem;Miktar\n01.09.2026;Servis;Güzergah 2 - Gebze;2\n", encoding="utf-8")
            s = main.calistir(ORNEK / "sozlesme.csv", t / "f.csv", t / "x.xlsx", t / "k.csv")
            turler = {u["tur"] for u in s["uyarilar"]}
            self.assertTrue({"Kullanım kaydı yok", "Faturası gelmemiş kullanım"} <= turler)

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx"), "--donem", "2026-09"]), 0)
            self.assertEqual(main.main(["--donem", "dün"]), 2)
            self.assertEqual(main.main(["--faturalar", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
