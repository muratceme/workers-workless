"""
Personel Maliyet Bütçesi — Workers / Workless kod bloğu
İnsan Kaynakları › İnsan Kaynakları Müdürü

Kadro planı, ücretler ve yan haklardan aylık personel maliyet bütçesi üretir; gerçekleşmeyle karşılaştırır:
  - Kadro: mevcut ve planlanan kadrolar; başlangıç ve bitiş ayı (yeni kadro, planlanan ayrılış).
  - Ücret artışı: --zam "7=10" → Temmuz'dan itibaren %10 (birden fazla artış birikimli uygulanır). Mevcut
    kadroya ve artış ayından önce başlayan yeni kadroya uygulanır.
  - Yan haklar: TL veya brüt ücretin %'si; kapsam (Herkes / Departman=… / Pozisyon~…), SGK'ya tabi mi, hangi aylar.
  - İşveren maliyeti = brüt ücret + yan haklar + SGK işveren payı + işsizlik sigortası işveren payı.
    Prime esas kazanç = brüt + SGK'ya tabi yan haklar; SGK tavanı ile sınırlı. Teşvik puanı (yok / imalat / diger)
    SGK işveren payından düşülür. Oranlar tr_parametreler.json'dan (yıl bazında) okunur.
  - Gerçekleşme (isteğe bağlı): departman × ay tutarı; sapma ve sapma oranı, eşiği aşanlar işaretlenir.
Rapor: aylık bütçe (departman × ay), kişi bazında, maliyet bileşenleri, gerçekleşme, varsayımlar, uyarılar.
İnternete bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: 14 kadro, 2026, Temmuz'da %10 artış
    python main.py --kadro kadro.xlsx --yan-haklar yan_haklar.csv --yil 2027 --zam 1=25 7=10
    python main.py --kadro kadro.xlsx --yan-haklar y.csv --yil 2026 --gerceklesen gerceklesen.xlsx --esik 5
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
PARAMETRE = BURASI / "tr_parametreler.json"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
K2 = Decimal("0.01")
SIFIR = Decimal(0)
AYLAR = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]

KADRO_SUTUNLARI = {"no": ("kadro no", "sicil no", "sicil", "no"), "pozisyon": ("pozisyon", "unvan", "ad soyad"), "departman": ("departman", "birim"),
                   "ucret": ("brut ucret tl", "brut ucret", "ucret", "aylik brut"), "bas": ("baslangic ayi", "baslangic", "giris ayi"),
                   "bit": ("bitis ayi", "bitis", "cikis ayi"), "tesvik": ("tesvik", "sgk tesvik")}
YAN_SUTUNLARI = {"kalem": ("kalem", "yan hak"), "tur": ("tur", "hesaplama"), "tutar": ("tutar", "deger", "oran"), "kapsam": ("kapsam",),
                 "sgk": ("sgk ya tabi", "sgk tabi", "sgk"), "aylar": ("aylar", "odeme aylari")}
GERCEK_SUTUNLARI = {"departman": ("departman", "birim"), "ay": ("ay", "donem"), "tutar": ("tutar tl", "tutar", "gerceklesen")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("%", "").replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def ay_oku(x) -> int | None:
    """'4', 'Nisan', '2026-04', '04.2026' → 4"""
    if x in (None, ""):
        return None
    s = katla(x)
    for i, ad in enumerate(AYLAR, 1):
        if s.startswith(katla(ad)):
            return i
    m = re.fullmatch(r"(\d{4}) (\d{1,2})|(\d{1,2}) (\d{4})|(\d{1,2})", s)
    if not m:
        return None
    v = int(m[2] or m[3] or m[5])
    return v if 1 <= v <= 12 else None


def tl(x: Decimal | None) -> str:
    return "—" if x is None else f"{x.quantize(K2, ROUND_HALF_UP):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


def parametreler(yil: int) -> tuple[dict, int, list[str]]:
    if not PARAMETRE.exists():
        raise ValueError("tr_parametreler.json bulunamadı")
    yillar = json.loads(PARAMETRE.read_text(encoding="utf-8")).get("yillar", {})
    for y in range(yil, yil - 3, -1):
        p = yillar.get(str(y), {})
        if p.get("sgk_isveren_orani") and p.get("sgk_tavan"):
            return p, y, ([] if y == yil else [f"{yil} SGK oranları ve tavanı tr_parametreler.json'da yok; {y} değerleri kullanıldı. "
                                              "Yeni dönem açıklanınca dosyayı güncelleyin."])
    raise ValueError(f"{yil} için SGK parametresi bulunamadı")


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Kadro:
    no: str
    pozisyon: str
    departman: str
    ucret: Decimal
    bas: int
    bit: int
    tesvik: str
    yeni: bool = False                               # başlangıç ayı verilmiş (yıl içinde başlayan kadro)
    aylik: dict = field(default_factory=dict)        # ay → {bileşen: tutar}

    def aktif(self, ay: int) -> bool:
        return self.bas <= ay <= self.bit

    def alan(self, ad: str) -> str:
        return {"departman": self.departman, "pozisyon": self.pozisyon, "kadro no": self.no}.get(ad, "")


@dataclass
class YanHak:
    kalem: str
    yuzde: bool
    tutar: Decimal
    kapsam: str
    sgk: bool
    aylar: set

    def uyar(self, k: Kadro) -> bool:
        s = self.kapsam.strip()
        if not s or katla(s) in ("herkes", "tumu"):
            return True
        m = re.fullmatch(r"\s*([^=~!]+?)\s*(!=|=|~)\s*(.+?)\s*", s)
        if not m:
            return False
        deger = katla(k.alan(katla(m[1])))
        hedef = katla(m[3])
        return deger == hedef if m[2] == "=" else deger != hedef if m[2] == "!=" else hedef in deger


def oku_kadro(yol: Path) -> tuple[list[Kadro], list[dict]]:
    sonuc, uy = [], []
    for r in kayitlar(yol, KADRO_SUTUNLARI, ("ucret",)):
        u = para(r.get("ucret"))
        if u is None:
            if r.get("no") or r.get("pozisyon"):
                uy.append({"onem": "Orta", "tur": "Okunamayan satır", "kim": metin(r.get("no")), "aciklama": f"Satır {r['_satir']}: ücret okunamadı"})
            continue
        b, e = ay_oku(r.get("bas")), ay_oku(r.get("bit"))
        if (r.get("bas") and b is None) or (r.get("bit") and e is None):
            uy.append({"onem": "Orta", "tur": "Ay okunamadı", "kim": metin(r.get("no")), "aciklama": f"Başlangıç / bitiş ayı anlaşılamadı: "
                       f"{metin(r.get('bas'))} / {metin(r.get('bit'))}; tüm yıl sayıldı"})
        t = katla(r.get("tesvik")) or "yok"
        sonuc.append(Kadro(metin(r.get("no")) or f"S{r['_satir']}", metin(r.get("pozisyon")), metin(r.get("departman")) or "—", u, b or 1, e or 12,
                           t if t in ("yok", "imalat", "diger") else "yok", b is not None))
    return sonuc, uy


def oku_yan_haklar(yol: Path | None) -> list[YanHak]:
    if not yol:
        return []
    y = []
    for r in kayitlar(yol, YAN_SUTUNLARI, ("kalem", "tutar")):
        t = para(r.get("tutar"))
        if t is None:
            continue
        aylar_s = metin(r.get("aylar"))
        aylar = set(range(1, 13)) if aylar_s in ("", "*") else {a for a in (ay_oku(x) for x in re.split(r"[,;/ ]+", aylar_s)) if a}
        y.append(YanHak(metin(r["kalem"]), "%" in metin(r.get("tur")) or katla(r.get("tur")).startswith(("yuzde", "oran")), t, metin(r.get("kapsam")),
                        katla(r.get("sgk")).startswith(("e", "1", "x")), aylar))
    return y


# ----------------------------------------------------------------------------
# Hesap
# ----------------------------------------------------------------------------

def hesapla(kadrolar: list[Kadro], yan: list[YanHak], p: dict, zam: dict[int, Decimal]) -> None:
    sgk = Decimal(p["sgk_isveren_orani"])
    iss = Decimal(p["issizlik_isveren_orani"])
    tavan = Decimal(p["sgk_tavan"])
    tesvik = {k: Decimal(v) for k, v in p.get("sgk_isveren_tesvik_puani", {}).items()}
    for k in kadrolar:
        for ay in range(1, 13):
            if not k.aktif(ay):
                continue
            carpan = Decimal(1)
            for zay, oran in sorted(zam.items()):
                if zay <= ay and (not k.yeni or k.bas < zay):
                    carpan *= 1 + oran
            ucret = (k.ucret * carpan).quantize(K2, ROUND_HALF_UP)
            x = {"Brüt ücret": ucret}
            tabi = ucret
            for y in yan:
                if ay in y.aylar and y.uyar(k):
                    t = (ucret * y.tutar / 100).quantize(K2, ROUND_HALF_UP) if y.yuzde else y.tutar
                    x[y.kalem] = x.get(y.kalem, SIFIR) + t
                    if y.sgk:
                        tabi += t
            matrah = min(tabi, tavan)
            x["SGK işveren payı"] = (matrah * (sgk - tesvik.get(k.tesvik, SIFIR))).quantize(K2, ROUND_HALF_UP)
            x["İşsizlik işveren payı"] = (matrah * iss).quantize(K2, ROUND_HALF_UP)
            x["_tavan"] = tabi > tavan
            k.aylik[ay] = x


def toplam(x: dict) -> Decimal:
    return sum((v for a, v in x.items() if not a.startswith("_")), SIFIR)


def karsilastir(kadrolar: list[Kadro], gercek: list[dict], esik: Decimal) -> tuple[list[dict], list[dict]]:
    butce = defaultdict(lambda: SIFIR)
    for k in kadrolar:
        for ay, x in k.aylik.items():
            butce[(katla(k.departman), ay)] += toplam(x)
    adlar = {katla(k.departman): k.departman for k in kadrolar}
    g = defaultdict(lambda: SIFIR)
    uy = []
    for r in gercek:
        ay, t = ay_oku(r.get("ay")), para(r.get("tutar"))
        if ay is None or t is None:
            continue
        d = katla(r.get("departman"))
        g[(d, ay)] += t
        adlar.setdefault(d, metin(r.get("departman")))
    if not g:
        return [], uy
    aylar = sorted({a for _, a in g})
    satirlar = []
    for d in sorted(adlar, key=lambda x: adlar[x]):
        for ay in aylar:
            b, gr = butce.get((d, ay), SIFIR), g.get((d, ay))
            if gr is None:
                if b:
                    uy.append({"onem": "Bilgi", "tur": "Gerçekleşen yok", "kim": adlar[d], "aciklama": f"{AYLAR[ay - 1]}: bütçe {tl(b)} TL, gerçekleşen kaydı yok"})
                continue
            sap = gr - b
            oran = sap / b if b else None
            satirlar.append({"departman": adlar[d], "ay": ay, "butce": b, "gercek": gr, "sapma": sap, "oran": oran})
            if b == 0:
                uy.append({"onem": "Orta", "tur": "Bütçesiz gerçekleşme", "kim": adlar[d], "aciklama": f"{AYLAR[ay - 1]}: {tl(gr)} TL gerçekleşme, bütçe yok"})
            elif abs(oran) > esik:
                uy.append({"onem": "Orta" if abs(oran) <= esik * 3 else "Yüksek", "tur": "Bütçe sapması", "kim": adlar[d],
                           "aciklama": f"{AYLAR[ay - 1]}: bütçe {tl(b)} TL, gerçekleşen {tl(gr)} TL (" + f"{oran * 100:+.1f}%".replace(".", ",") + ")"})
    return satirlar, uy


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
UST = Alignment(vertical="top", wrap_text=True)
PF = "#,##0"


def _baslik(ws, basliklar, genislik=()):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "B2"


def rapor_yaz(cikti: Path, s: dict) -> None:
    kadrolar = s["kadrolar"]
    wb = Workbook()
    ab = wb.active
    ab.title = "Aylık Bütçe"
    _baslik(ab, ["Departman"] + AYLAR + ["Yıllık"], [18] + [12] * 13)
    deps = sorted({k.departman for k in kadrolar})
    for d in deps:
        v = [sum((toplam(k.aylik[a]) for k in kadrolar if k.departman == d and a in k.aylik), SIFIR) for a in range(1, 13)]
        ab.append([d] + [float(x) for x in v] + [float(sum(v))])
    v = [sum((toplam(k.aylik[a]) for k in kadrolar if a in k.aylik), SIFIR) for a in range(1, 13)]
    ab.append(["Toplam"] + [float(x) for x in v] + [float(sum(v))])
    for h in ab[ab.max_row]:
        h.font = Font(bold=True)
    ab.append(["Kişi sayısı"] + [sum(1 for k in kadrolar if k.aktif(a)) for a in range(1, 13)] + [None])
    for row in ab.iter_rows(min_row=2, max_row=ab.max_row - 1, min_col=2):
        for c in row:
            c.number_format = PF
    n = len(deps)
    g = BarChart()
    g.type, g.grouping, g.overlap = "col", "stacked", 100
    g.title, g.height, g.width = "Aylık personel maliyeti (departman)", 8, 22
    g.add_data(Reference(ab, min_col=1, max_col=13, min_row=2, max_row=1 + n), titles_from_data=True, from_rows=True)
    g.set_categories(Reference(ab, min_col=2, max_col=13, min_row=1))
    ab.add_chart(g, f"A{ab.max_row + 3}")

    bl = wb.create_sheet("Maliyet Bileşenleri")
    bilesenler = list(dict.fromkeys(a for k in kadrolar for x in k.aylik.values() for a in x if not a.startswith("_")))
    _baslik(bl, ["Bileşen"] + AYLAR + ["Yıllık", "Pay"], [26] + [12] * 13 + [8])
    genel = sum((toplam(x) for k in kadrolar for x in k.aylik.values()), SIFIR)
    for b in bilesenler:
        v = [sum((k.aylik[a].get(b, SIFIR) for k in kadrolar if a in k.aylik), SIFIR) for a in range(1, 13)]
        bl.append([b] + [float(x) for x in v] + [float(sum(v)), float(sum(v) / genel) if genel else None])
        for j in range(2, 15):
            bl.cell(bl.max_row, j).number_format = PF
        bl.cell(bl.max_row, 15).number_format = "0.0%"

    kb = wb.create_sheet("Kişi Bazında")
    _baslik(kb, ["Kadro No", "Pozisyon", "Departman", "Başlangıç", "Bitiş", "Teşvik", "Brüt (başlangıç)", "Brüt (Aralık)", "Yıllık Brüt",
                 "Yıllık Yan Haklar", "Yıllık SGK + İşsizlik", "Yıllık İşveren Maliyeti", "Aylık Ortalama", "SGK Tavanı"],
            (9, 26, 14, 9, 9, 8, 13, 13, 14, 14, 14, 15, 13, 9))
    for k in kadrolar:
        xs = list(k.aylik.values())
        brut = sum((x["Brüt ücret"] for x in xs), SIFIR)
        sgk = sum((x["SGK işveren payı"] + x["İşsizlik işveren payı"] for x in xs), SIFIR)
        top = sum((toplam(x) for x in xs), SIFIR)
        son = k.aylik.get(max(k.aylik), {}).get("Brüt ücret") if k.aylik else None
        kb.append([k.no, k.pozisyon, k.departman, AYLAR[k.bas - 1], AYLAR[k.bit - 1], k.tesvik, float(k.ucret), float(son) if son else None, float(brut),
                   float(top - brut - sgk), float(sgk), float(top), float(top / len(xs)) if xs else None, "Evet" if any(x["_tavan"] for x in xs) else ""])
        for j in range(7, 14):
            kb.cell(kb.max_row, j).number_format = PF
    kb.auto_filter.ref = f"A1:N{kb.max_row}"

    if s["karsilastirma"]:
        gk = wb.create_sheet("Gerçekleşme")
        _baslik(gk, ["Departman", "Ay", "Bütçe", "Gerçekleşen", "Sapma", "Sapma %", "Açıklama"], (18, 10, 14, 14, 13, 9, 40))
        for x in s["karsilastirma"]:
            gk.append([x["departman"], AYLAR[x["ay"] - 1], float(x["butce"]), float(x["gercek"]), float(x["sapma"]),
                       float(x["oran"]) if x["oran"] is not None else None, ""])
            for j in (3, 4, 5):
                gk.cell(gk.max_row, j).number_format = PF
            gk.cell(gk.max_row, 6).number_format = "0.0%"
            if x["oran"] is not None and abs(x["oran"]) > s["esik"]:
                gk.cell(gk.max_row, 6).fill = PatternFill("solid", fgColor="FDE2E1" if x["oran"] > 0 else "E3F4E1")
            gk.cell(gk.max_row, 7).fill = PatternFill("solid", fgColor="FFF4CE")
        tb = sum((x["butce"] for x in s["karsilastirma"]), SIFIR)
        tg = sum((x["gercek"] for x in s["karsilastirma"]), SIFIR)
        gk.append(["Toplam", "", float(tb), float(tg), float(tg - tb), float((tg - tb) / tb) if tb else None])
        for j in (3, 4, 5):
            gk.cell(gk.max_row, j).number_format = PF
        gk.cell(gk.max_row, 6).number_format = "0.0%"

    vs = wb.create_sheet("Varsayımlar")
    _baslik(vs, ["Varsayım", "Değer"], (40, 80))
    p = s["p"]
    for a, v in [("Bütçe yılı", s["yil"]), ("Parametre yılı (tr_parametreler.json)", s["param_yili"]),
                 ("SGK işveren payı", f"%{Decimal(p['sgk_isveren_orani']) * 100:g}"), ("İşsizlik sigortası işveren payı", f"%{Decimal(p['issizlik_isveren_orani']) * 100:g}"),
                 ("SGK tavanı (aylık)", f"{tl(Decimal(p['sgk_tavan']))} TL"),
                 ("Teşvik puanları", ", ".join(f"{k}: {Decimal(v) * 100:g} puan" for k, v in p.get("sgk_isveren_tesvik_puani", {}).items())),
                 ("Ücret artışları", ", ".join(f"{AYLAR[a - 1]} %{o * 100:g}" for a, o in sorted(s["zam"].items())) or "yok"),
                 ("Yan haklar", "; ".join(f"{y.kalem} ({'%' + format(y.tutar, 'g') + ' brüt' if y.yuzde else tl(y.tutar) + ' TL'}, {y.kapsam or 'Herkes'}, "
                                          f"SGK {'tabi' if y.sgk else 'tabi değil'}, aylar: {'tümü' if len(y.aylar) == 12 else ','.join(map(str, sorted(y.aylar)))})"
                                          for y in s["yan"]) or "yok"),
                 ("Kapsam dışı", "Kıdem tazminatı karşılığı, fazla mesai, ihbar ve izin karşılığı, kıst ay hesabı ve damga vergisi dahil değildir; "
                                 "giriş ve çıkış ayları tam ay sayılır. SGK tavanını aşan ikramiye tutarının sonraki aylara devri hesaplanmaz.")]:
        vs.append([a, v])
        vs.cell(vs.max_row, 2).alignment = UST

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Kim", "Açıklama"], (9, 22, 16, 90))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["kim"], u["aciklama"]])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(kadro_yolu: Path, cikti: Path, yil: int, yan_yolu: Path | None = None, gercek_yolu: Path | None = None,
             zam: dict[int, Decimal] | None = None, esik: Decimal = Decimal("0.05")) -> dict:
    kadrolar, uyarilar = oku_kadro(kadro_yolu)
    if not kadrolar:
        raise ValueError(f"{kadro_yolu.name}: kadro bulunamadı")
    p, param_yili, notlar = parametreler(yil)
    uyarilar += [{"onem": "Orta", "tur": "Parametre", "kim": "", "aciklama": n} for n in notlar]
    yan = oku_yan_haklar(yan_yolu)
    zam = zam or {}
    hesapla(kadrolar, yan, p, zam)
    if p.get("asgari_ucret_brut") and param_yili == yil:
        asg = Decimal(p["asgari_ucret_brut"])
        for k in kadrolar:
            if k.ucret < asg:
                uyarilar.append({"onem": "Yüksek", "tur": "Asgari ücret altı", "kim": k.no, "aciklama": f"Brüt {tl(k.ucret)} TL < {yil} asgari ücreti {tl(asg)} TL"})
    for k in kadrolar:
        if any(x["_tavan"] for x in k.aylik.values()):
            uyarilar.append({"onem": "Bilgi", "tur": "SGK tavanı", "kim": k.no, "aciklama": "Prime esas kazanç bazı aylarda SGK tavanını aşıyor; SGK payı tavandan hesaplandı"})
        if k.bas > k.bit:
            uyarilar.append({"onem": "Orta", "tur": "Ay hatası", "kim": k.no, "aciklama": "Başlangıç ayı bitiş ayından sonra; kadro bütçeye girmedi"})
    gercek = kayitlar(gercek_yolu, GERCEK_SUTUNLARI, ("departman", "ay", "tutar")) if gercek_yolu else []
    kars, u2 = karsilastir(kadrolar, gercek, esik)
    uyarilar += u2
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    uyarilar.sort(key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))
    s = {"kadrolar": kadrolar, "yan": yan, "p": p, "yil": yil, "param_yili": param_yili, "zam": zam, "karsilastirma": kars, "uyarilar": uyarilar,
         "esik": esik, "toplam": sum((toplam(x) for k in kadrolar for x in k.aylik.values()), SIFIR)}
    rapor_yaz(cikti, s)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Kadro planı, ücretler ve yan haklardan aylık personel maliyet bütçesi üretir, gerçekleşmeyle karşılaştırır.")
    p.add_argument("--kadro", type=Path, default=ORNEK / "kadro.csv", help="Kadro No, Pozisyon, Departman, Brüt Ücret, Başlangıç Ayı, Bitiş Ayı, Teşvik")
    p.add_argument("--yan-haklar", type=Path, help="Kalem, Tür (TL / %%), Tutar, Kapsam, SGK'ya Tabi, Aylar")
    p.add_argument("--gerceklesen", type=Path, help="İsteğe bağlı: Departman, Ay, Tutar (TL)")
    p.add_argument("--yil", type=int, help="Bütçe yılı (örnek veride 2026)")
    p.add_argument("--zam", nargs="*", default=None, metavar="AY=ORAN", help="Ücret artışı, ör. 1=25 7=10 (Ocak %%25, Temmuz %%10)")
    p.add_argument("--esik", type=float, default=5, help="Sapma eşiği %% (varsayılan 5)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "personel_maliyet_butcesi.xlsx")
    a = p.parse_args(argv)
    ornek = a.kadro == ORNEK / "kadro.csv"
    yan = a.yan_haklar or (ORNEK / "yan_haklar.csv" if ornek else None)
    ger = a.gerceklesen or (ORNEK / "gerceklesen.csv" if ornek else None)
    for y in (a.kadro, yan, ger):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    zam_arg = a.zam if a.zam is not None else (["7=10"] if ornek else [])
    zam = {}
    for x in zam_arg:
        ay, _, o = x.partition("=")
        if ay_oku(ay) is None or para(o) is None:
            print(f"[X] --zam 'AY=ORAN' biçiminde olmalı (ör. 7=10): {x}")
            return 2
        zam[ay_oku(ay)] = para(o) / 100
    yil = a.yil or (2026 if ornek else None)
    if not yil:
        print("[X] --yil verin (ör. --yil 2027)")
        return 2
    try:
        s = calistir(a.kadro, a.cikti, yil, yan, ger, zam, Decimal(str(a.esik)) / 100)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(s['kadrolar'])} kadro · {yil} yıllık işveren maliyeti {tl(s['toplam'])} TL")
    if s["karsilastirma"]:
        tb = sum((x["butce"] for x in s["karsilastirma"]), SIFIR)
        tg = sum((x["gercek"] for x in s["karsilastirma"]), SIFIR)
        print(f"[OK] Gerçekleşme ({len({x['ay'] for x in s['karsilastirma']})} ay): bütçe {tl(tb)} TL, gerçekleşen {tl(tg)} TL, sapma {tl(tg - tb)} TL")
    for u in s["uyarilar"]:
        if u["onem"] != "Bilgi":
            print(f"[!] {u['kim']} {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
