import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "p.xlsx"
        cls.s = main.calistir(ORNEK / "kadro.csv", cls.cikti, 2026, ORNEK / "yan_haklar.csv", ORNEK / "gerceklesen.csv", {7: D("0.10")})
        cls.k = {k.no: k for k in cls.s["kadrolar"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_aylik_maliyet(self):
        x = self.k["K01"].aylik[1]          # 145.000 + yemek 4.200 + yol 2.500 + ÖSS 1.800 + araç 32.000; SGK (147.500) × (%21,75 − 5) + işsizlik %2
        self.assertEqual((x["SGK işveren payı"], x["İşsizlik işveren payı"]), (D("24706.25"), D("2950.00")))
        self.assertEqual(main.toplam(x), D("213156.25"))
        x = self.k["K13"].aylik[1]          # tavan 297.270
        self.assertEqual(x["SGK işveren payı"], D("64656.23"))
        self.assertTrue(x["_tavan"])

    def test_zam_ikramiye_kadro(self):
        self.assertEqual(self.k["K02"].aylik[7]["Brüt ücret"], D("79200.00"))
        self.assertEqual(self.k["K02"].aylik[6]["İkramiye"], D("72000.00"))
        self.assertEqual(self.k["K02"].aylik[12]["İkramiye"], D("79200.00"))
        self.assertEqual(sorted(self.k["K07"].aylik), list(range(4, 13)))
        self.assertEqual(sorted(self.k["K10"].aylik), list(range(1, 9)))
        self.assertEqual(self.k["K14"].aylik[9]["Brüt ücret"], D("54000"))       # Temmuz artışından sonra başladı
        self.assertNotIn("Yol yardımı", self.k["K08"].aylik[1])
        self.assertIn("Araç tahsisi (kira bedeli)", self.k["K13"].aylik[1])       # Pozisyon~Müdür
        self.assertEqual(self.s["toplam"], D("21390936.49"))

    def test_gerceklesme(self):
        kars = {(x["departman"], x["ay"]): x for x in self.s["karsilastirma"]}
        self.assertAlmostEqual(float(kars[("Üretim", 8)]["oran"]), 0.09, places=4)
        u = [(x["tur"], x["kim"]) for x in self.s["uyarilar"]]
        self.assertEqual(sorted(k for t, k in u if t == "Bütçe sapması"), ["Satış", "Üretim"])
        self.assertEqual(sum(1 for t, k in u if t == "Gerçekleşen yok" and k == "İnsan Kaynakları"), 9)

    def test_excel(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Aylık Bütçe", "Maliyet Bileşenleri", "Kişi Bazında", "Gerçekleşme", "Varsayımlar", "Uyarılar"])
        self.assertEqual(wb["Aylık Bütçe"]["N7"].value, 21390936.49)


class KuralTesti(unittest.TestCase):
    def test_yardimcilar(self):
        self.assertEqual([main.ay_oku(x) for x in ("4", "Nisan", "2027-04", "04.2027", "13", "")], [4, 4, 4, 4, None, None])
        k = main.Kadro("1", "Satış Müdürü", "Satış", D(1), 1, 12, "yok")
        self.assertTrue(main.YanHak("x", False, D(1), "Departman=satış", False, set()).uyar(k))
        self.assertFalse(main.YanHak("x", False, D(1), "Departman!=Satış", False, set()).uyar(k))
        self.assertTrue(main.YanHak("x", False, D(1), "Pozisyon~müdür", False, set()).uyar(k))

    def test_parametre_yedegi_ve_asgari(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t / "k.csv").write_text("Kadro No;Departman;Brüt Ücret;Başlangıç Ayı\nA;X;30000;\nB;X;30000;1\n", encoding="utf-8")
            s = main.calistir(t / "k.csv", t / "o.xlsx", 2027)
            self.assertEqual(s["param_yili"], 2026)
            self.assertTrue(any(u["tur"] == "Parametre" for u in s["uyarilar"]))
            s = main.calistir(t / "k.csv", t / "o.xlsx", 2026, zam={1: D("0.1")})
            self.assertTrue(any(u["tur"] == "Asgari ücret altı" for u in s["uyarilar"]))
            self.assertEqual(s["kadrolar"][0].aylik[1]["Brüt ücret"], D("33000.00"))   # mevcut kadro Ocak artışını alır
            self.assertEqual(s["kadrolar"][1].aylik[1]["Brüt ücret"], D("30000"))      # Ocak'ta başlayan yeni kadro almaz

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--zam", "Temmuz"]), 2)
            self.assertEqual(main.main(["--kadro", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
