import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
RAPOR = date(2026, 10, 16)
TATIL = {date(2026, 10, 29)}


def hedef(*satirlar):
    """(py, kalem, hedef, ağırlık)"""
    return {(main.katla(p), main.katla(k)): {"py": p, "kalem": k, "hedef": D(str(h)), "birim": "TL", "agirlik": D(str(a)) if a is not None else None} for p, k, h, a in satirlar}


def gercek(*satirlar):
    return {"deger": {(main.katla(p), main.katla(k)): D(str(v)) for p, k, v in satirlar}, "ad": {(main.katla(p), main.katla(k)): (p, k) for p, k, v in satirlar}}


class HesapTesti(unittest.TestCase):
    def test_is_gunleri(self):
        g = main.is_gunleri(2026, 10, TATIL)
        self.assertEqual((len(g), sum(d <= RAPOR for d in g)), (21, 12))            # Ekim 2026: 22 hafta içi − 29 Ekim
        self.assertEqual(len(main.is_gunleri(2026, 10, set())), 22)
        self.assertEqual(main.ay_coz("2026-10"), (2026, 10))
        self.assertEqual(main.ay_coz("10.2026"), (2026, 10))
        self.assertIsNone(main.ay_coz("Ekim"))

    def test_temel_formuller(self):
        r = main.hesapla(hedef(("A", "Mevduat", 2100, 100)), gercek(("A", "Mevduat", 1050)), RAPOR, TATIL)
        s = r["satirlar"][0]
        self.assertEqual(r["gun"], (12, 21, 9))
        self.assertEqual((s["oran"], s["olmasi"], s["sapma"], s["tahmin"], s["gunluk"], s["durum"]),
                         (D("50.0"), D("1200.00"), D("-150.00"), D("1837.50"), D("116.67"), "Geride"))   # 1837,5 < 2100 × 0,9 = 1890

    def test_durumlar(self):
        def d(g):
            return main.hesapla(hedef(("A", "K", 2100, 100)), gercek(("A", "K", g)), RAPOR, TATIL)["satirlar"][0]["durum"]
        self.assertEqual([d(2100), d(1200), d(1199), d(1080), d(1079)], ["Hedef tamam", "Yolunda", "Risk", "Risk", "Geride"])

    def test_karne_puani(self):
        h = hedef(("A", "K1", 100, 60), ("A", "K2", 100, 40), ("B", "K1", 100, 60), ("B", "K2", 100, 40))
        g = gercek(("A", "K1", 200), ("A", "K2", 50), ("B", "K1", -30), ("B", "K2", 90))
        r = main.hesapla(h, g, RAPOR, TATIL)
        self.assertEqual({k["py"]: (k["sira"], k["puan"]) for k in r["karne"]}, {"A": (1, D("92.0")), "B": (2, D("36.0"))})   # 0,6×120 + 0,4×50; eksi → 0
        self.assertEqual(main.hesapla(h, g, RAPOR, TATIL, D(200))["karne"][0]["puan"], D("140.0"))

    def test_agirliksiz_ve_sube(self):
        h = hedef(("A", "K1", 100, None), ("A", "K2", 300, None), ("B", "K1", 300, None))
        r = main.hesapla(h, gercek(("A", "K1", 80), ("A", "K2", 60), ("B", "K1", 120), ("B", "Fon", 10)), RAPOR, TATIL)
        self.assertEqual(r["karne"][0]["puan"], D("50.0"))                           # ağırlık yoksa eşit: (80 + 20) / 2
        s = r["sube"][0]
        self.assertEqual((s["hedef"], s["gercek"], s["oran"], s["en_iyi"], s["en_zayif"]), (D(400), D(200), D("50.0"), "A (%80.0)", "B (%40.0)"))
        self.assertEqual([x["durum"] for x in r["satirlar"] if x["kalem"] == "Fon"], ["Hedef yok"])
        self.assertEqual(r["kalemler"], ["K1", "K2"])

    def test_agirlik_uyarilari(self):
        r = main.hesapla(hedef(("A", "K1", 100, 50), ("A", "K2", 100, None), ("B", "K1", 100, 40), ("B", "K2", 100, 40)), gercek(), RAPOR, TATIL)
        self.assertEqual(sorted(u["tur"] for u in r["uyarilar"]), ["Ağırlık toplamı 100 değil", "Ağırlığı olmayan kalem"])

    def test_ay_basi(self):
        with self.assertRaises(ValueError):
            main.hesapla(hedef(("A", "K", 1, 1)), gercek(), date(2026, 8, 1), set())      # 1 Ağustos 2026 Cumartesi


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "p.xlsx"
            r = main.calistir(ORNEK / "hedefler.csv", ORNEK / "gerceklesmeler.csv", cikti, RAPOR, ORNEK / "tatiller.csv")
            self.assertEqual((len(r["karne"]), len(r["kalemler"]), r["gun"]), (5, 6, (12, 21, 9)))
            self.assertEqual(r["karne"][0]["py"], "A. Yılmaz")
            tek = next(s for s in r["satirlar"] if (s["py"], s["kalem"]) == ("E. Tekin", "Ticari Kredi"))
            self.assertEqual((tek["gercek"], tek["durum"]), (D(0), "Geride"))
            self.assertEqual(sorted(u["tur"] for u in r["uyarilar"]), ["Dönem dışı satır", "Hedefi olmayan gerçekleşme"])
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Pano", "Detay", "Şube Toplamı", "Uyarılar"])
            self.assertEqual(wb["Pano"].cell(3, 2).value, "A. Yılmaz")

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            h, g = Path(tmp) / "h.csv", Path(tmp) / "g.csv"
            h.write_text("Ay;Portföy Yöneticisi;Kalem;Hedef\n2026-10;A;K;1.000,00\n2026-10;A;K;500\n2026-11;A;Z;100\n2026-10;B;K;sıfır\n", encoding="utf-8")
            g.write_text("Tarih;Portföy Yöneticisi;Kalem;Gerçekleşen\n05.10.2026;a;k;250,50\n30.09.2026;A;K;999\nbozuk;A;K;5\n06.10.2026;A;K;149,50\n", encoding="utf-8")
            hd, gr, _, uy = main.oku(h, g, RAPOR)
            self.assertEqual((list(hd), gr["deger"][("a", "k")]), ([("a", "k")], D("400.00")))
            self.assertEqual(sorted(u["tur"] for u in uy), ["Dönem dışı satır", "Gerçekleşme satırı okunamadı", "Hedef okunamadı", "Mükerrer hedef"])
            arg = ["--hedefler", str(h), "--gerceklesmeler", str(g), "--cikti", str(Path(tmp) / "x.xlsx")]
            self.assertEqual(main.main(arg + ["--tarih", "16.10.2026"]), 0)
            self.assertEqual(main.main(arg + ["--tarih", "16.12.2026"]), 1)                # o ay için hedef yok
            self.assertEqual(main.main(arg + ["--tarih", "yarın"]), 2)
            self.assertEqual(main.main(["--hedefler", str(Path(tmp) / "yok.csv")]), 1)
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "y.xlsx")]), 0)


if __name__ == "__main__":
    unittest.main()
