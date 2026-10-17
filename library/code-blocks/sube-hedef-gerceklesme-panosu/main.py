"""
Şube Hedef-Gerçekleşme Panosu — Gerisi Bende kod bloğu
Şube Bankacılığı › Şube Müdürü

Mevduat, kredi, kart, sigorta gibi kalemlerin aylık hedeflerini portföy yöneticisi bazında gerçekleşmelerle karşılaştırır
ve tek panoda toplar.
  Gerçekleşme %        = gerçekleşen ÷ hedef
  Olması gereken       = hedef × geçen iş günü ÷ ayın toplam iş günü
  Ay sonu tahmini      = gerçekleşen ÷ geçen iş günü × toplam iş günü   (aynı tempo sürerse)
  Günlük gereken       = (hedef − gerçekleşen) ÷ kalan iş günü
  Karne puanı          = Σ ağırlık × min(gerçekleşme %, tavan) ÷ Σ ağırlık   (tavan varsayılan %120; eksi gerçekleşme 0 sayılır)
  Durum: Hedef tamam (gerçekleşen ≥ hedef) · Yolunda (tahmin ≥ hedef) · Risk (tahmin ≥ hedefin %90'ı) · Geride.
İş günü: Pazartesi–Cuma; resmî tatiller --tatiller dosyasıyla düşülür (yarım günler tam gün sayılır).
Gerçekleşme satırları rapor tarihine kadar (dahil) toplanır; net azalışlar eksi tutarla girilebilir.
İnternete bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: Ekim 2026, 5 portföy yöneticisi, 6 kalem
    python main.py --hedefler hedefler.xlsx --gerceklesmeler g.xlsx --tarih 16.10.2026 --tatiller tatiller.csv
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


HEDEF_SUTUNLARI = {"py": ("portfoy yoneticisi", "musteri temsilcisi", "personel", "py"), "kalem": ("kalem", "urun", "hedef kalemi"), "hedef": ("hedef", "aylik hedef", "donem hedefi"),
                   "birim": ("birim",), "agirlik": ("agirlik", "agirlik yuzde", "karne agirligi"), "ay": ("ay", "donem")}
GERCEK_SUTUNLARI = {"tarih": ("tarih", "islem tarihi"), "py": ("portfoy yoneticisi", "musteri temsilcisi", "personel", "py"), "kalem": ("kalem", "urun", "hedef kalemi"),
                    "tutar": ("gerceklesen", "tutar", "adet", "miktar")}
DURUM_RENK = {"Hedef tamam": YESIL, "Yolunda": YESIL, "Risk": SARI, "Geride": KIRMIZI, "Hedef yok": MAVI}


def is_gunleri(yil: int, ay: int, tatiller: set[date]) -> list[date]:
    d, l = date(yil, ay, 1), []
    while d.month == ay:
        if d.weekday() < 5 and d not in tatiller:
            l.append(d)
        d += timedelta(days=1)
    return l


def ay_coz(x) -> tuple[int, int] | None:
    """'2026-10', '10.2026', '10/2026' veya tarih → (yıl, ay)"""
    t = tarih(x)
    if t:
        return t.year, t.month
    m = re.fullmatch(r"\s*(\d{4})[-./](\d{1,2})\s*", metin(x)) or None
    if m:
        return int(m.group(1)), int(m.group(2))
    m = re.fullmatch(r"\s*(\d{1,2})[-./](\d{4})\s*", metin(x))
    return (int(m.group(2)), int(m.group(1))) if m else None


def oku(h_yol: Path, g_yol: Path, rapor: date, t_yol: Path | None = None) -> tuple[dict, dict, set, list[dict]]:
    uy, hedefler, gercek, tatiller = [], {}, defaultdict(Decimal), set()
    if t_yol:
        for r in tablo_oku(t_yol):
            t = tarih(r[0])
            if t:
                tatiller.add(t)
    for r in kayitlar(h_yol, HEDEF_SUTUNLARI, ("py", "kalem", "hedef")):
        py, kalem = metin(r["py"]), metin(r["kalem"])
        if not py or not kalem:
            continue
        if metin(r["ay"]) and ay_coz(r["ay"]) != (rapor.year, rapor.month):
            continue                                                   # başka ayın hedefi
        h = para(r["hedef"])
        if h is None or h <= 0:
            uy.append(uyari("Yüksek", "Hedef okunamadı", f"{py} · {kalem}", f"Satır {r['_satir']}: hedef '{metin(r['hedef'])}'. Sıfırdan büyük olmalı; satır atlandı."))
            continue
        a = para(r["agirlik"])
        if (katla(py), katla(kalem)) in hedefler:
            uy.append(uyari("Orta", "Mükerrer hedef", f"{py} · {kalem}", f"Satır {r['_satir']}: aynı kişi ve kalem için ikinci hedef; ilk satır kullanıldı."))
            continue
        hedefler[(katla(py), katla(kalem))] = {"py": py, "kalem": kalem, "hedef": h, "birim": metin(r["birim"]), "agirlik": a if a is not None and a > 0 else None}
    if not hedefler:
        raise ValueError(f"{h_yol.name}: {rapor:%m.%Y} dönemi için hedef bulunamadı")
    adlar, disinda = {}, 0
    for r in kayitlar(g_yol, GERCEK_SUTUNLARI, ("py", "kalem", "tutar")):
        py, kalem = metin(r["py"]), metin(r["kalem"])
        if not py or not kalem:
            continue
        t, v = tarih(r["tarih"]), para(r["tutar"])
        if v is None or (metin(r["tarih"]) and not t):
            uy.append(uyari("Yüksek", "Gerçekleşme satırı okunamadı", f"{py} · {kalem}", f"Satır {r['_satir']}: tarih '{metin(r['tarih'])}', tutar '{metin(r['tutar'])}'. Satır atlandı."))
            continue
        if t and ((t.year, t.month) != (rapor.year, rapor.month) or t > rapor):
            disinda += 1
            continue
        gercek[(katla(py), katla(kalem))] += v
        adlar[(katla(py), katla(kalem))] = (py, kalem)
    if disinda:
        uy.append(uyari("Bilgi", "Dönem dışı satır", g_yol.name, f"{disinda} gerçekleşme satırı {rapor:%m.%Y} dönemi dışında veya rapor tarihinden sonra; toplama katılmadı."))
    for k, v in gercek.items():
        if k not in hedefler:
            uy.append(uyari("Orta", "Hedefi olmayan gerçekleşme", f"{adlar[k][0]} · {adlar[k][1]}", f"{v:f} gerçekleşme var, hedef tanımlı değil. Panoda 'Hedef yok' olarak gösterildi."))
    return hedefler, {"deger": gercek, "ad": adlar}, tatiller, uy


def hesapla(hedefler: dict, gercek: dict, rapor: date, tatiller: set[date] = frozenset(), tavan: Decimal = Decimal(120)) -> dict:
    gunler = is_gunleri(rapor.year, rapor.month, set(tatiller))
    toplam, gecen = len(gunler), sum(d <= rapor for d in gunler)
    kalan = toplam - gecen
    uy, satirlar = [], []
    if not gecen:
        raise ValueError(f"{rapor:%d.%m.%Y}: ayın ilk iş günü henüz geçmemiş; tahmin yapılamaz")

    def satir(py, kalem, hedef, g, birim="", agirlik=None):
        if hedef is None:
            return {"py": py, "kalem": kalem, "hedef": None, "gercek": g, "oran": None, "olmasi": None, "sapma": None, "tahmin": yuvarla(g / gecen * toplam), "gunluk": None,
                    "durum": "Hedef yok", "birim": birim, "agirlik": None}
        tahmin = g / gecen * toplam
        durum = "Hedef tamam" if g >= hedef else "Yolunda" if tahmin >= hedef else "Risk" if tahmin >= hedef * Decimal("0.9") else "Geride"
        olmasi = hedef * gecen / toplam
        return {"py": py, "kalem": kalem, "hedef": hedef, "gercek": g, "oran": yuvarla(g / hedef * 100, "0.1"), "olmasi": yuvarla(olmasi), "sapma": yuvarla(g - olmasi),
                "tahmin": yuvarla(tahmin), "gunluk": yuvarla(max(hedef - g, Decimal(0)) / kalan) if kalan else None, "durum": durum, "birim": birim, "agirlik": agirlik}

    for k, h in hedefler.items():
        satirlar.append(satir(h["py"], h["kalem"], h["hedef"], gercek["deger"].get(k, Decimal(0)), h["birim"], h["agirlik"]))
    for k, v in gercek["deger"].items():
        if k not in hedefler:
            satirlar.append(satir(*gercek["ad"][k], None, v))
    kisiler, kalemler = list(dict.fromkeys(s["py"] for s in satirlar)), list(dict.fromkeys(s["kalem"] for s in satirlar if s["hedef"] is not None))
    karne = []
    for py in kisiler:
        l = [s for s in satirlar if s["py"] == py and s["hedef"] is not None]
        if not l:
            continue
        agirliksiz = [s["kalem"] for s in l if s["agirlik"] is None]
        if agirliksiz and len(agirliksiz) < len(l):
            uy.append(uyari("Orta", "Ağırlığı olmayan kalem", py, f"{', '.join(agirliksiz)} için ağırlık girilmemiş; karne puanına katılmadı."))
        w = [(s["agirlik"] if s["agirlik"] is not None else (Decimal(1) if len(agirliksiz) == len(l) else Decimal(0)), s) for s in l]
        wt = sum(a for a, _ in w)
        if not agirliksiz and wt != 100:
            uy.append(uyari("Bilgi", "Ağırlık toplamı 100 değil", py, f"Ağırlıklar toplamı {wt:f}; puan ağırlık toplamına bölünerek hesaplandı."))
        puan = sum(a * min(max(s["oran"], Decimal(0)), tavan) for a, s in w) / wt if wt else None
        karne.append({"py": py, "puan": yuvarla(puan, "0.1") if puan is not None else None, "tamam": sum(s["durum"] == "Hedef tamam" for s in l),
                      "geride": sum(s["durum"] == "Geride" for s in l), "kalem": len(l)})
    karne.sort(key=lambda k: -(k["puan"] if k["puan"] is not None else -1))
    for i, k in enumerate(karne, 1):
        k["sira"] = i
    sube = []
    for kalem in kalemler:
        l = [s for s in satirlar if s["kalem"] == kalem and s["hedef"] is not None]
        t = satir("Şube", kalem, sum(s["hedef"] for s in l), sum(s["gercek"] for s in l), l[0]["birim"])
        en_iyi, en_zayif = max(l, key=lambda s: s["oran"]), min(l, key=lambda s: s["oran"])
        sube.append(t | {"kisi": len(l), "en_iyi": f"{en_iyi['py']} (%{en_iyi['oran']})", "en_zayif": f"{en_zayif['py']} (%{en_zayif['oran']})"})
    return {"satirlar": satirlar, "karne": karne, "sube": sube, "kalemler": kalemler, "uyarilar": uy, "rapor": rapor, "gun": (gecen, toplam, kalan), "tavan": tavan}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

def _oran_renk(oran, beklenen) -> str | None:
    if oran is None:
        return None
    return YESIL if oran >= beklenen else SARI if oran >= beklenen * Decimal("0.9") else KIRMIZI


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    gecen, toplam, kalan = r["gun"]
    beklenen = Decimal(gecen) / toplam * 100
    pn = wb.active
    pn.title = "Pano"
    pn.append([f"Rapor tarihi {r['rapor']:%d.%m.%Y} · geçen iş günü {gecen} / {toplam} · beklenen gerçekleşme %{yuvarla(beklenen, '0.1')}"])
    pn["A1"].font = Font(bold=True)
    _baslik(pn, ["Sıra", "Portföy Yöneticisi", "Karne Puanı", *[f"{k} %" for k in r["kalemler"]], "Tamamlanan", "Geride"], (6, 24, 11, *[14] * len(r["kalemler"]), 11, 8), "C3")
    hucre = {(s["py"], s["kalem"]): s for s in r["satirlar"]}
    for k in r["karne"]:
        oranlar = [hucre[(k["py"], kl)]["oran"] if (k["py"], kl) in hucre else None for kl in r["kalemler"]]
        _satir(pn, [k["sira"], k["py"], k["puan"], *oranlar, f"{k['tamam']} / {k['kalem']}", k["geride"]],
               dolgu={3: _oran_renk(k["puan"], beklenen)} | {4 + i: _oran_renk(o, beklenen) for i, o in enumerate(oranlar)})
    _satir(pn, ["", "Şube", "", *[s["oran"] for s in r["sube"]], "", ""], dolgu={4 + i: _oran_renk(s["oran"], beklenen) for i, s in enumerate(r["sube"])})
    for c in pn[pn.max_row]:
        c.font = Font(bold=True)
    pn.append([])
    pn.append([f"Renkler ayın geçen bölümüne göre: yeşil = beklenen orana ulaşmış, sarı = beklenenin %90'ı ve üzeri, kırmızı = altında. Karne puanında kalem başına tavan %{r['tavan']:f}."])

    dt = wb.create_sheet("Detay")
    bas = ["Portföy Yöneticisi", "Kalem", "Birim", "Hedef", "Gerçekleşen", "Gerçekleşme %", "Olması Gereken", "Sapma", "Ay Sonu Tahmini", "Kalan Günlük Gereken", "Durum", "Ağırlık"]
    _baslik(dt, bas, (24, 22, 8, 16, 16, 11, 16, 16, 16, 16, 13, 9))
    for s in r["satirlar"]:
        _satir(dt, [s["py"], s["kalem"], s["birim"], s["hedef"], s["gercek"], s["oran"], s["olmasi"], s["sapma"], s["tahmin"], s["gunluk"], s["durum"], s["agirlik"]],
               para_s=(4, 5, 7, 8, 9, 10), dolgu={11: DURUM_RENK[s["durum"]]})
    dt.auto_filter.ref = f"A1:L{dt.max_row}"

    sb = wb.create_sheet("Şube Toplamı")
    _baslik(sb, ["Kalem", "Birim", "Hedef", "Gerçekleşen", "Gerçekleşme %", "Olması Gereken", "Sapma", "Ay Sonu Tahmini", "Kalan Günlük Gereken", "Durum", "En İyi", "En Zayıf"],
            (22, 8, 16, 16, 11, 16, 16, 16, 16, 13, 28, 28))
    for s in r["sube"]:
        _satir(sb, [s["kalem"], s["birim"], s["hedef"], s["gercek"], s["oran"], s["olmasi"], s["sapma"], s["tahmin"], s["gunluk"], s["durum"], s["en_iyi"], s["en_zayif"]],
               para_s=(3, 4, 6, 7, 8, 9), dolgu={10: DURUM_RENK[s["durum"]]})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(h_yol: Path, g_yol: Path, cikti: Path, rapor: date, t_yol: Path | None = None, tavan: Decimal = Decimal(120)) -> dict:
    hedefler, gercek, tatiller, uy = oku(h_yol, g_yol, rapor, t_yol)
    r = hesapla(hedefler, gercek, rapor, tatiller, tavan)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Şube hedeflerinin portföy yöneticisi bazında gerçekleşmesini tek panoda toplar.")
    p.add_argument("--hedefler", type=Path, default=ORNEK / "hedefler.csv", help="Portföy Yöneticisi, Kalem, Hedef, Birim, Ağırlık, Ay (isteğe bağlı: 2026-10)")
    p.add_argument("--gerceklesmeler", type=Path, default=ORNEK / "gerceklesmeler.csv", help="Tarih, Portföy Yöneticisi, Kalem, Gerçekleşen")
    p.add_argument("--tatiller", type=Path, help="Her satırda bir resmî tatil tarihi (GG.AA.YYYY)")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 16.10.2026)")
    p.add_argument("--tavan", type=Decimal, default=Decimal(120), help="Karne puanında kalem başına en yüksek gerçekleşme %% (varsayılan 120)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "hedef_gerceklesme_panosu.xlsx")
    a = p.parse_args(argv)
    ornek = a.hedefler == ORNEK / "hedefler.csv"
    tat = a.tatiller or (ORNEK / "tatiller.csv" if ornek else None)
    for y in (a.hedefler, a.gerceklesmeler, tat):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    rapor = tarih(a.tarih) if a.tarih else (date(2026, 10, 16) if ornek else date.today())
    if not rapor or a.tavan <= 0:
        print("[X] --tarih GG.AA.YYYY biçiminde, --tavan sıfırdan büyük olmalı")
        return 2
    try:
        r = calistir(a.hedefler, a.gerceklesmeler, a.cikti, rapor, tat, a.tavan)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    gecen, toplam, _ = r["gun"]
    print(f"[OK] {len(r['karne'])} portföy yöneticisi · {len(r['kalemler'])} kalem · iş günü {gecen}/{toplam}")
    for s in r["sube"]:
        print(f"     {s['kalem']}: %{s['oran']} · {s['durum']}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['kim']} · {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
