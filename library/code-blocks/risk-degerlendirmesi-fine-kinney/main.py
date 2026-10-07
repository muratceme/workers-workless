"""
Risk Değerlendirmesi (Fine-Kinney) — Workers / Workless kod bloğu
İş Sağlığı ve Güvenliği › İş Güvenliği Uzmanı

Tehlike listesindeki olasılık (O), frekans (F) ve şiddet (Ş) puanlarından Fine-Kinney risk skorunu
(R = O × F × Ş) ve risk sınıfını hesaplar; geçersiz puanları, önlemi eksik yüksek riskleri ve
terminini geçmiş aksiyonları işaretler; önlem sonrası (kalıntı) riski karşılaştırır. Sonuç, sınıfa
göre renklendirilmiş ve önceliğe göre sıralanmış bir Excel risk değerlendirme tablosudur.

Skala ve sınıflar: Kinney ve Wiruth (1976), Türkçe uygulama — bkz. README.

Kullanım:
    python main.py                                   # örnek tehlike listesiyle dener
    python main.py --girdi tehlike_listesi.xlsx --cikti cikti/risk_degerlendirmesi.xlsx
"""
from __future__ import annotations

import argparse
import csv
import sys
from collections import Counter
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent

OLASILIK = {Decimal("10"): "Kuvvetle beklenir", Decimal("6"): "Oldukça mümkün", Decimal("3"): "Olağan dışı fakat olabilir",
            Decimal("1"): "Çok uzak ihtimal", Decimal("0.5"): "İhtimal dahilinde fakat beklenmez",
            Decimal("0.2"): "Pratik olarak imkânsız", Decimal("0.1"): "Neredeyse imkânsız"}
FREKANS = {Decimal("10"): "Sürekli (saatlik)", Decimal("6"): "Sıklıkla (günlük)", Decimal("3"): "Ara sıra (haftalık)",
           Decimal("2"): "Nadir (aylık)", Decimal("1"): "Seyrek (yıllık)", Decimal("0.5"): "Oldukça seyrek (yılda belki bir)"}
SIDDET = {Decimal("100"): "Facia (birden fazla ölüm)", Decimal("40"): "Felaket (ölümlü kaza)",
          Decimal("15"): "Çok ciddi (yaralanma / iş günü kaybı)", Decimal("7"): "Ciddi (yaralanma, dış ilk yardım)",
          Decimal("3"): "Önemli (dahili ilk yardım)", Decimal("1"): "Fark edilebilir (ucuz atlatma)"}

# (alt sınır dahil, sınıf, eylem, renk)
SINIFLAR = [
    (Decimal(400), "Çok Yüksek Risk", "Tolerans gösterilemez: faaliyet durdurulmalı, hemen önlem alınmalı", "C00000"),
    (Decimal(200), "Yüksek Risk", "Kısa vadeli eylem planına alınmalı", "FF6B6B"),
    (Decimal(70), "Önemli Risk", "Dikkatle izlenmeli, eylem planına alınmalı", "FFB347"),
    (Decimal(20), "Olası Risk", "Eylem planına alınmalı, gözetim altında tutulmalı", "FFE08A"),
    (Decimal(0), "Kabul Edilebilir Risk", "Acil eylem gerekmeyebilir", "B7E4C7"),
]


def sayi(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x)).normalize()
    try:
        return Decimal(str(x).strip().replace(",", ".")).normalize()
    except Exception:
        return None


def sinif(r: Decimal) -> tuple[str, str, str]:
    for alt, ad, eylem, renk in SINIFLAR:
        if r >= alt:
            return ad, eylem, renk
    return SINIFLAR[-1][1:]


def puan_kontrol(deger, skala: dict, ad: str) -> tuple[Decimal | None, str]:
    d = sayi(deger)
    if d is None:
        return None, f"{ad} boş"
    if d not in skala:
        izinli = ", ".join(f"{k:g}" for k in sorted(skala))
        return None, f"{ad} {d:g} skalada yok (izinli: {izinli})"
    return d, ""


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


SUTUNLAR = {
    "bolum": ("bölüm", "bölüm / faaliyet", "faaliyet", "alan", "birim"),
    "tehlike": ("tehlike", "tehlike kaynağı"),
    "risk": ("risk", "olası sonuç"),
    "o": ("olasılık", "o"),
    "f": ("frekans", "f"),
    "s": ("şiddet", "ş", "s"),
    "mevcut": ("mevcut önlemler", "mevcut önlem"),
    "onerilen": ("önerilen önlemler", "alınacak önlemler", "düzeltici faaliyet", "önlem"),
    "sorumlu": ("sorumlu",),
    "termin": ("termin", "termin tarihi", "son tarih"),
    "o2": ("önlem sonrası olasılık", "kalıntı olasılık", "o2"),
    "f2": ("önlem sonrası frekans", "kalıntı frekans", "f2"),
    "s2": ("önlem sonrası şiddet", "kalıntı şiddet", "ş2", "s2"),
}


def tarih(x) -> date | None:
    if x in (None, ""):
        return None
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(str(x).strip(), f).date()
        except ValueError:
            pass
    return None


def liste_oku(yol: Path) -> list[dict]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ayirici = csv.Sniffer().sniff(metin.splitlines()[0], delimiters=",;\t").delimiter
        satirlar = list(csv.reader(metin.splitlines(), delimiter=ayirici))
    satirlar = [r for r in satirlar if any(c not in (None, "") for c in r)]
    b = [kucuk(x) for x in satirlar[0]]
    k = {alan: next((i for i, x in enumerate(b) if x in adlar), None) for alan, adlar in SUTUNLAR.items()}
    eksik = [a for a in ("tehlike", "o", "f", "s") if k[a] is None]
    if eksik:
        raise SystemExit(f"Gerekli sütunlar bulunamadı: Tehlike, Olasılık, Frekans, Şiddet. Başlıklar: {satirlar[0]}")
    return [{alan: (r[i] if i is not None and i < len(r) else None) for alan, i in k.items()} for r in satirlar[1:]]


def degerlendir(kayit: dict, bugun: date) -> dict:
    uyarilar = []
    o, h1 = puan_kontrol(kayit["o"], OLASILIK, "Olasılık")
    f, h2 = puan_kontrol(kayit["f"], FREKANS, "Frekans")
    s, h3 = puan_kontrol(kayit["s"], SIDDET, "Şiddet")
    uyarilar += [h for h in (h1, h2, h3) if h]
    sonuc = dict(kayit)
    if None in (o, f, s):
        sonuc.update(r=None, sinif="Hesaplanamadı", eylem="", renk="DDDDDD")
    else:
        r = (o * f * s).normalize()
        ad, eylem, renk = sinif(r)
        sonuc.update(r=r, sinif=ad, eylem=eylem, renk=renk)
        if r >= 70 and not str(kayit.get("onerilen") or "").strip():
            uyarilar.append("Önemli ve üzeri risk için önerilen önlem yazılmamış")
        if r >= 70 and not str(kayit.get("sorumlu") or "").strip():
            uyarilar.append("Sorumlu atanmamış")
    t = tarih(kayit.get("termin"))
    if t and t < bugun and sonuc.get("r") is not None and sonuc["r"] >= 20:
        uyarilar.append(f"Termin geçti ({t:%d.%m.%Y})")
    # Önlem sonrası (kalıntı) risk
    if any(kayit.get(x) not in (None, "") for x in ("o2", "f2", "s2")):
        o2, g1 = puan_kontrol(kayit["o2"], OLASILIK, "Önlem sonrası olasılık")
        f2, g2 = puan_kontrol(kayit["f2"], FREKANS, "Önlem sonrası frekans")
        s2, g3 = puan_kontrol(kayit["s2"], SIDDET, "Önlem sonrası şiddet")
        uyarilar += [h for h in (g1, g2, g3) if h]
        if None not in (o2, f2, s2):
            r2 = (o2 * f2 * s2).normalize()
            sonuc.update(r2=r2, sinif2=sinif(r2)[0])
            if sonuc.get("r") is not None and r2 > sonuc["r"]:
                uyarilar.append("Önlem sonrası risk, önlem öncesinden yüksek")
    sonuc["uyarilar"] = "; ".join(uyarilar)
    return sonuc


BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")


def rapor_yaz(sonuclar: list[dict], cikti: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Risk Değerlendirmesi"
    kolonlar = [("No", None), ("Bölüm / Faaliyet", "bolum"), ("Tehlike", "tehlike"), ("Risk", "risk"),
                ("O", "o"), ("F", "f"), ("Ş", "s"), ("Risk Skoru (R)", "r"), ("Risk Sınıfı", "sinif"), ("Eylem", "eylem"),
                ("Mevcut Önlemler", "mevcut"), ("Önerilen Önlemler", "onerilen"), ("Sorumlu", "sorumlu"), ("Termin", "termin"),
                ("Önlem Sonrası R", "r2"), ("Önlem Sonrası Sınıf", "sinif2"), ("Uyarılar", "uyarilar")]
    ws.append([k for k, _ in kolonlar])
    for h in ws[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    sirali = sorted(sonuclar, key=lambda x: -(x.get("r") or -1))
    for n, s in enumerate(sirali, 1):
        satir = []
        for _, alan in kolonlar:
            v = n if alan is None else s.get(alan)
            if isinstance(v, Decimal):
                v = float(v)
            satir.append(v)
        ws.append(satir)
        ws.cell(ws.max_row, 9).fill = PatternFill("solid", fgColor=s["renk"])
        ws.cell(ws.max_row, 8).fill = PatternFill("solid", fgColor=s["renk"])
        for h in ws[ws.max_row]:
            h.alignment = Alignment(wrap_text=True, vertical="top")
    genislik = [5, 18, 30, 26, 6, 6, 6, 10, 18, 34, 30, 34, 14, 12, 10, 18, 40]
    for j, g in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = g
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions

    oz = wb.create_sheet("Özet")
    oz.append(["Risk Sınıfı", "Adet", "Eylem"])
    for h in oz[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    sayac = Counter(s["sinif"] for s in sonuclar)
    for _, ad, eylem, renk in SINIFLAR:
        oz.append([ad, sayac.get(ad, 0), eylem])
        oz.cell(oz.max_row, 1).fill = PatternFill("solid", fgColor=renk)
    if sayac.get("Hesaplanamadı"):
        oz.append(["Hesaplanamadı", sayac["Hesaplanamadı"], "Puanları skaladaki değerlerle düzeltin"])
    oz.append([])
    oz.append(["Bölüm", "Tehlike Sayısı", "En Yüksek R"])
    for h in oz[oz.max_row]:
        h.font = Font(bold=True)
    bolumler = {}
    for s in sonuclar:
        b = s.get("bolum") or "-"
        x = bolumler.setdefault(b, [0, Decimal(0)])
        x[0] += 1
        x[1] = max(x[1], s.get("r") or Decimal(0))
    for b, (adet, ust) in sorted(bolumler.items(), key=lambda i: -i[1][1]):
        oz.append([b, adet, float(ust)])
    for c, g in zip("ABC", (24, 14, 60)):
        oz.column_dimensions[c].width = g

    sk = wb.create_sheet("Skalalar")
    for baslik, skala in (("Olasılık (O)", OLASILIK), ("Frekans (F)", FREKANS), ("Şiddet (Ş)", SIDDET)):
        sk.append([baslik])
        sk.cell(sk.max_row, 1).font = Font(bold=True)
        for v, ac in sorted(skala.items(), reverse=True):
            sk.append([float(v), ac])
        sk.append([])
    sk.append(["Risk sınıfları (R = O × F × Ş)"])
    sk.cell(sk.max_row, 1).font = Font(bold=True)
    for alt, ad, eylem, _ in SINIFLAR:
        sk.append([f"R ≥ {alt}" if alt else "R < 20", f"{ad}: {eylem}"])
    sk.append([])
    sk.append(["Kaynak", "Kinney G.F., Wiruth A.D. (1976), Practical Risk Analysis for Safety Management; Türkçe uygulama skalaları"])
    sk.column_dimensions["A"].width = 22
    sk.column_dimensions["B"].width = 90
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(girdi: Path, cikti: Path, bugun: date | None = None) -> list[dict]:
    sonuclar = [degerlendir(k, bugun or date.today()) for k in liste_oku(girdi)]
    rapor_yaz(sonuclar, cikti)
    return sonuclar


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description="Fine-Kinney risk değerlendirmesi.")
    ap.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "tehlike_listesi.csv")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "risk_degerlendirmesi.xlsx")
    a = ap.parse_args(argv)
    sonuclar = calistir(a.girdi, a.cikti)
    sayac = Counter(s["sinif"] for s in sonuclar)
    print(f"[OK] {len(sonuclar)} tehlike değerlendirildi: " + ", ".join(f"{k} {v}" for k, v in sayac.most_common()))
    print(f"[{'!' if any(s['uyarilar'] for s in sonuclar) else 'OK'}] {sum(1 for s in sonuclar if s['uyarilar'])} satırda uyarı var")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
