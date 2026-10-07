import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))

import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402


class SinifTesti(unittest.TestCase):
    def test_esikler(self):
        cases = [(D(19), "Kabul Edilebilir Risk"), (D(20), "Olası Risk"), (D("69.9"), "Olası Risk"), (D(70), "Önemli Risk"),
                 (D(199), "Önemli Risk"), (D(200), "Yüksek Risk"), (D(399), "Yüksek Risk"), (D(400), "Çok Yüksek Risk")]
        for r, beklenen in cases:
            self.assertEqual(main.sinif(r)[0], beklenen, r)

    def test_skala_disi_puan(self):
        s = main.degerlendir({"tehlike": "x", "o": 3, "f": 6, "s": 5, "onerilen": "a", "sorumlu": "b"}, date(2026, 1, 1))
        self.assertIsNone(s["r"])
        self.assertIn("Şiddet 5 skalada yok", s["uyarilar"])

    def test_ondalik_virgul(self):
        s = main.degerlendir({"tehlike": "x", "o": "0,5", "f": "6", "s": "40", "onerilen": "a", "sorumlu": "b"}, date(2026, 1, 1))
        self.assertEqual(s["r"], D(120))


class UctanUcaTest(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "r.xlsx"
            sonuclar = main.calistir(KLASOR / "ornek_veri" / "tehlike_listesi.csv", cikti, date(2026, 10, 7))
            by = {s["tehlike"]: s for s in sonuclar}
            forklift = by["Forkliftin yaya yolunu kullanması"]
            self.assertEqual((forklift["r"], forklift["sinif"]), (D(720), "Çok Yüksek Risk"))
            self.assertEqual((forklift["r2"], forklift["sinif2"]), (D(120), "Önemli Risk"))
            pano = by["Pano kapağının açık bırakılması"]
            self.assertEqual(pano["r"], D(120))
            self.assertIn("önlem yazılmamış", pano["uyarilar"])
            self.assertIn("Termin geçti", by["Kaygan zemin (yağ sızıntısı)"]["uyarilar"])
            self.assertEqual(by["Gürültü (88 dB)"]["sinif"], "Hesaplanamadı")
            ws = load_workbook(cikti)["Risk Değerlendirmesi"]
            self.assertEqual(ws["H2"].value, 720)          # en yüksek risk ilk sırada


if __name__ == "__main__":
    unittest.main()
