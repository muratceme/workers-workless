"""
Geri Ödeme Kapasitesi Analizi — Workers / Workless kod bloğu
Bankacılık › Krediler (Tahsis) › Kredi Tahsis Uzmanı

Firmanın nakit yaratma gücünü mevcut ve önerilen kredilerin borç servisiyle karşılaştırır:
  - Geçmiş mali tablolardan varsayımlar: satış büyümesi (bileşik), FAVÖK marjı, vergi/FAVÖK, yatırım/satış,
    net işletme sermayesi (ticari alacak + stok − ticari borç)/satış.
  - Projeksiyon: borç servisine kullanılabilir nakit (CFADS) = FAVÖK − vergi − yatırım − NİS artışı.
  - Borç servisi: mevcut krediler (eşit taksit, eşit anapara, vade sonu, rotatif = yalnız faiz) ve önerilen
    kredi (ödemesiz dönemli) aylık ödeme planlarından yıllık anapara + faiz; dövizli krediler kurla TL'ye.
  - DSCR = CFADS / borç servisi (yıl bazında) ve net finansal borç / FAVÖK.
  - Senaryolar: satış düşüşü, marj daralması, değişken faiz artışı, kur şoku ve birleşik stres.
  - Asgari DSCR'yi (--min-dscr) baz senaryoda sağlayan azami kredi tutarı ve DSCR'yi 1,0'a indiren satış
    düşüşü (kırılma noktası).
İnternete bağlanmaz. Kredi kararı ve iç derecelendirme bankanın yetkili organlarına aittir.

Kullanım:
    python main.py                                                     # örnek firma ve 20 milyon TL kredi önerisi
    python main.py --mali mali_tablolar.xlsx --borclar borclar.xlsx --oneri kredi_onerisi.json --min-dscr 1,2
    python main.py --mali mali.xlsx --borclar borclar.xlsx --oneri oneri.json --buyume 25
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import re
import sys
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
TL_ADLARI = {"", "tl", "try"}
KALEMLER = {
    "satis": ("net satislar", "net satis", "hasilat", "satis gelirleri"),
    "favok": ("favok", "ebitda"),
    "vergi": ("odenen vergi", "vergi", "vergi odemesi", "donem kari vergi gideri"),
    "yatirim": ("yatirim harcamalari", "yatirim", "capex", "maddi duran varlik alimi"),
    "alacak": ("ticari alacaklar", "ticari alacak"),
    "stok": ("stoklar", "stok"),
    "borc": ("ticari borclar", "ticari borc"),
    "hazir": ("hazir degerler", "nakit ve nakit benzerleri", "nakit"),
}
VARSAYILAN_SENARYOLAR = [
    {"ad": "Baz", "satis": 0, "marj": 0, "faiz": 0, "kur": 0},
    {"ad": "Satış −%15", "satis": -15, "marj": 0, "faiz": 0, "kur": 0},
    {"ad": "FAVÖK marjı −3 puan", "satis": 0, "marj": -3, "faiz": 0, "kur": 0},
    {"ad": "Değişken faiz +10 puan", "satis": 0, "marj": 0, "faiz": 10, "kur": 0},
    {"ad": "Kur +%30", "satis": 0, "marj": 0, "faiz": 0, "kur": 30},
    {"ad": "Birleşik stres", "satis": -10, "marj": -2, "faiz": 5, "kur": 20},
]


def katla(s) -> str:
    return " ".join(str(s or "").translate(_TR).lower().split())


def sayi(x) -> float | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return float(x)
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1 or (s.count(".") == 1 and len(s.split(".")[1]) == 3):
        s = s.replace(".", "")
    try:
        return float(Decimal(s))
    except InvalidOperation:
        return None


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


# ----------------------------------------------------------------------------
# Girdiler
# ----------------------------------------------------------------------------

def mali_oku(yol: Path) -> tuple[list[str], dict[str, dict[str, float]]]:
    s = tablo_oku(yol)
    yillar = [str(x).strip() for x in s[0][1:] if str(x or "").strip()]
    veri = {y: {} for y in yillar}
    for r in s[1:]:
        k = katla(r[0])
        alan = next((a for a, es in KALEMLER.items() if k in es), None)
        if alan:
            for y, v in zip(yillar, r[1:]):
                if sayi(v) is not None:
                    veri[y][alan] = sayi(v)
    eksik = [a for a in ("satis", "favok") if any(a not in veri[y] for y in yillar)]
    if eksik:
        raise SystemExit(f"{yol.name}: Net Satışlar ve FAVÖK her yıl için gerekli (eksik: {', '.join(eksik)})")
    return yillar, veri


@dataclass
class Kredi:
    ad: str
    bakiye: float
    para: str
    faiz: float               # yıllık %
    vade: int                 # ay
    tip: str                  # esit-taksit | esit-anapara | vade-sonu | rotatif
    degisken: bool
    odemesiz: int = 0
    yeni: bool = False
    notlar: list[str] = field(default_factory=list)


def tip_coz(x) -> str:
    k = katla(x)
    if "rotatif" in k or "spot" in k or "kmh" in k or "yalniz faiz" in k:
        return "rotatif"
    if "anapara" in k and "esit" in k:
        return "esit-anapara"
    if "vade sonu" in k or "bullet" in k or "tek odeme" in k:
        return "vade-sonu"
    return "esit-taksit"


def borclari_oku(yol: Path) -> list[Kredi]:
    s = tablo_oku(yol)
    b = [katla(x) for x in s[0]]
    bul = lambda *a: next((i for i, x in enumerate(b) if x in a), None)  # noqa: E731
    k = {"ad": bul("kredi", "kredi adi", "urun"), "banka": bul("banka"), "bakiye": bul("bakiye", "risk", "anapara bakiyesi"),
         "pb": bul("para birimi", "doviz", "pb"), "faiz": bul("yillik faiz (%)", "yillik faiz", "faiz", "faiz orani"),
         "vade": bul("kalan vade (ay)", "kalan vade", "vade (ay)"), "tip": bul("odeme tipi", "geri odeme", "odeme sekli"),
         "faiz_tipi": bul("faiz tipi", "faiz turu")}
    if k["bakiye"] is None:
        raise SystemExit(f"{yol.name}: 'Bakiye' sütunu yok. Başlıklar: {s[0]}")
    al = lambda r, a: r[k[a]] if k[a] is not None and k[a] < len(r) else None  # noqa: E731
    sonuc = []
    for i, r in enumerate(s[1:], 1):
        bakiye = sayi(al(r, "bakiye"))
        if not bakiye:
            continue
        ad = " · ".join(str(x).strip() for x in (al(r, "ad"), al(r, "banka")) if x) or f"Kredi {i}"
        tip = tip_coz(al(r, "tip"))
        vade = int(sayi(al(r, "vade")) or 12)
        kr = Kredi(ad, bakiye, str(al(r, "pb") or "TL").strip().upper(), sayi(al(r, "faiz")) or 0.0, vade, tip,
                   "degisken" in katla(al(r, "faiz_tipi")) or tip == "rotatif")
        if tip == "rotatif":
            kr.notlar.append("Rotatif: yalnız faiz servisi; anaparanın vadede yenileneceği varsayıldı")
        if sayi(al(r, "vade")) is None:
            kr.notlar.append("Kalan vade yok; 12 ay varsayıldı")
        sonuc.append(kr)
    return sonuc


def oneri_oku(yol: Path) -> tuple[Kredi, dict[str, float]]:
    o = json.loads(yol.read_text(encoding="utf-8"))
    kr = Kredi("Önerilen kredi", float(o["tutar"]), str(o.get("para_birimi", "TL")).upper(), float(o["yillik_faiz"]), int(o["vade_ay"]),
               tip_coz(o.get("odeme_tipi", "Eşit taksit")), "degisken" in katla(o.get("faiz_tipi", "")), int(o.get("odemesiz_ay", 0)), yeni=True)
    return kr, {k.upper(): float(v) for k, v in o.get("kur", {}).items()}


# ----------------------------------------------------------------------------
# Ödeme planı ve projeksiyon
# ----------------------------------------------------------------------------

def odeme_plani(tutar: float, yillik_faiz: float, vade: int, tip: str, odemesiz: int = 0) -> list[tuple[float, float, float]]:
    """Aylık (anapara, faiz, dönem sonu bakiye). Ödemesiz dönemde faiz ödenir; aylık faiz = yıllık / 12."""
    r = yillik_faiz / 100 / 12
    bakiye, plan = tutar, []
    n = max(vade - odemesiz, 1)
    taksit = (tutar / n if r == 0 else tutar * r / (1 - (1 + r) ** -n)) if tip == "esit-taksit" else 0.0
    for ay in range(1, vade + 1):
        faiz = bakiye * r
        if tip == "rotatif" or ay <= odemesiz:
            anapara = 0.0
        elif tip == "vade-sonu":
            anapara = bakiye if ay == vade else 0.0
        elif tip == "esit-anapara":
            anapara = tutar / n
        else:
            anapara = taksit - faiz
        anapara = min(anapara, bakiye)
        bakiye -= anapara
        plan.append((anapara, faiz, bakiye))
    return plan


def yillik_servis(krediler: list[Kredi], yil_sayisi: int, kurlar: dict[str, float], faiz_soku: float = 0, kur_soku: float = 0
                  ) -> tuple[dict[str, list[tuple[float, float, float]]], list[str]]:
    """Kredi → yıl bazında (anapara, faiz, yıl sonu bakiye), TL."""
    sonuc, uyarilar = {}, []
    for k in krediler:
        kur = 1.0 if katla(k.para) in TL_ADLARI else kurlar.get(k.para)
        if kur is None:
            uyarilar.append(f"{k.ad}: {k.para} kuru yok; kredi hesaba katılmadı (--kur {k.para}=...)")
            continue
        kur *= 1 + kur_soku / 100
        faiz = k.faiz + (faiz_soku if k.degisken else 0)
        plan = odeme_plani(k.bakiye, faiz, k.vade, k.tip, k.odemesiz)
        yillar = []
        for y in range(yil_sayisi):
            ay = plan[12 * y:12 * (y + 1)]
            if k.tip == "rotatif" and y * 12 >= k.vade:     # yenilendiği varsayılan rotatif: faiz servisi sürer
                ay = [(0.0, k.bakiye * faiz / 1200, k.bakiye)] * 12
            bakiye = ay[-1][2] if ay else 0.0
            yillar.append((sum(a for a, _, _ in ay) * kur, sum(f for _, f, _ in ay) * kur, bakiye * kur))
        sonuc[k.ad] = yillar
    return sonuc, uyarilar


def varsayimlar(yillar: list[str], veri: dict, buyume: float | None) -> dict:
    s = [veri[y]["satis"] for y in yillar]
    nis = lambda y: veri[y].get("alacak", 0) + veri[y].get("stok", 0) - veri[y].get("borc", 0)  # noqa: E731
    son = yillar[-1]
    cagr = (s[-1] / s[0]) ** (1 / (len(s) - 1)) - 1 if len(s) > 1 and s[0] > 0 else 0.0
    toplam = lambda a: sum(veri[y].get(a, 0) for y in yillar)  # noqa: E731
    return {"satis0": s[-1], "marj": veri[son]["favok"] / s[-1], "buyume": cagr if buyume is None else buyume / 100, "buyume_kaynak":
            "geçmiş bileşik büyüme" if buyume is None else "kullanıcı", "cagr": cagr,
            "vergi_orani": toplam("vergi") / toplam("favok") if toplam("favok") else 0.0,
            "yatirim_orani": toplam("yatirim") / toplam("satis"), "nis_orani": nis(son) / s[-1], "nis0": nis(son),
            "hazir": veri[son].get("hazir", 0.0), "favok0": veri[son]["favok"]}


def projeksiyon(v: dict, servis: dict, yil_sayisi: int, satis_soku: float = 0, marj_soku: float = 0) -> list[dict]:
    satirlar, nis_onceki = [], v["nis0"]
    for y in range(1, yil_sayisi + 1):
        satis = v["satis0"] * (1 + v["buyume"]) ** y * (1 + satis_soku / 100)
        favok = satis * (v["marj"] + marj_soku / 100)
        vergi = max(favok, 0) * v["vergi_orani"]
        yatirim = satis * v["yatirim_orani"]
        nis = satis * v["nis_orani"]
        dnis = nis - nis_onceki
        cfads = favok - vergi - yatirim - dnis
        nis_onceki = nis
        anapara = sum(x[y - 1][0] for x in servis.values())
        faiz = sum(x[y - 1][1] for x in servis.values())
        borc = sum(x[y - 1][2] for x in servis.values())
        ds = anapara + faiz
        satirlar.append({"yil": y, "satis": satis, "favok": favok, "vergi": vergi, "yatirim": yatirim, "nis_degisim": dnis,
                         "cfads": cfads, "anapara": anapara, "faiz": faiz, "servis": ds, "dscr": cfads / ds if ds else None,
                         "borc": borc, "netborc_favok": (borc - v["hazir"]) / favok if favok > 0 else None})
    return satirlar


def min_dscr(p: list[dict]) -> float | None:
    d = [x["dscr"] for x in p if x["dscr"] is not None]
    return min(d) if d else None


def senaryo_calistir(v, krediler, kurlar, yil_sayisi, s) -> list[dict]:
    servis, _ = yillik_servis(krediler, yil_sayisi, kurlar, s["faiz"], s["kur"])
    return projeksiyon(v, servis, yil_sayisi, s["satis"], s["marj"])


def azami_kredi(v, mevcut, oneri, kurlar, yil_sayisi, esik) -> float:
    def uygun(tutar):
        k = Kredi(oneri.ad, tutar, oneri.para, oneri.faiz, oneri.vade, oneri.tip, oneri.degisken, oneri.odemesiz, True)
        m = min_dscr(senaryo_calistir(v, mevcut + ([k] if tutar > 0 else []), kurlar, yil_sayisi, VARSAYILAN_SENARYOLAR[0]))
        return m is None or m >= esik
    if not uygun(0):
        return 0.0
    alt, ust = 0.0, max(v["favok0"] * 20, oneri.bakiye * 4)
    for _ in range(50):
        orta = (alt + ust) / 2
        alt, ust = (orta, ust) if uygun(orta) else (alt, orta)
    return math.floor(alt / 1000) * 1000


def kirilma_satis(v, krediler, kurlar, yil_sayisi) -> float | None:
    """Baz senaryoda en düşük DSCR'yi 1,0'a indiren satış düşüşü (%)."""
    f = lambda s: min_dscr(senaryo_calistir(v, krediler, kurlar, yil_sayisi, {"satis": -s, "marj": 0, "faiz": 0, "kur": 0}))  # noqa: E731
    if f(0) is None or f(0) < 1:
        return 0.0
    alt, ust = 0.0, 95.0
    if f(ust) >= 1:
        return None
    for _ in range(50):
        orta = (alt + ust) / 2
        alt, ust = (orta, ust) if f(orta) >= 1 else (alt, orta)
    return round(alt, 1)


def sonuc_etiketi(m: float | None, esik: float) -> str:
    if m is None:
        return "Borç servisi yok"
    return "Karşılıyor" if m >= esik else "Sınırda" if m >= 1 else "Karşılamıyor"


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Karşılıyor": "E3F5E1", "Sınırda": "FFF4CE", "Karşılamıyor": "FDE2E1", "Borç servisi yok": "E8F0FE"}
UST = Alignment(vertical="top", wrap_text=True)
TL_BICIM = "#,##0"


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def tl(x: float) -> str:
    return f"{x:,.0f}".replace(",", ".")


def rapor_yaz(cikti: Path, s: dict) -> None:
    v, esik, oneri = s["varsayim"], s["esik"], s["oneri"]
    wb = Workbook()
    o = wb.active
    o.title = "Özet"
    o.column_dimensions["A"].width, o.column_dimensions["B"].width, o.column_dimensions["C"].width = 40, 22, 60
    baz = s["senaryolar"]["Baz"]
    satirlar = [
        ("Önerilen kredi", f"{tl(oneri.bakiye)} {oneri.para}", f"%{oneri.faiz:g} yıllık, {oneri.vade} ay, {oneri.odemesiz} ay ödemesiz, {oneri.tip}"),
        ("Asgari DSCR (eşik)", esik, "Bankanızın kredi politikasındaki eşiği kullanın; değer örnektir"),
        ("Baz senaryo en düşük DSCR", round(min_dscr(baz.projeksiyon), 2) if min_dscr(baz.projeksiyon) else "-", baz.sonuc),
        ("Baz senaryoda taşınabilir azami kredi (aynı yapıyla)", s["azami"], "Bu tutarın üstünde en düşük DSCR eşiğin altına iner"),
        ("DSCR'yi 1,0'a indiren satış düşüşü", "yok (%95'e kadar)" if s["kirilma"] is None else f"%{s['kirilma']:g}".replace(".", ","),
         "Baz senaryoya göre, tüm projeksiyon yıllarında"),
        ("İlk yıl net finansal borç / FAVÖK", round(baz.projeksiyon[0]["netborc_favok"], 2) if baz.projeksiyon[0]["netborc_favok"] else "-", ""),
    ]
    for a, b, c in satirlar:
        o.append([a, b, c])
        o.cell(o.max_row, 1).font = Font(bold=True)
    o.cell(5, 2).number_format = TL_BICIM
    o.cell(3, 3).fill = PatternFill("solid", fgColor=RENK.get(baz.sonuc, "FFFFFF"))
    o.append([])
    _baslik(o, ["Senaryo", "En düşük DSCR", "Sonuç"], ())
    for ad, x in s["senaryolar"].items():
        m = min_dscr(x.projeksiyon)
        o.append([ad, None if m is None else round(m, 2), x.sonuc])
        o.cell(o.max_row, 3).fill = PatternFill("solid", fgColor=RENK[x.sonuc])
    o.append([])
    for u in s["uyarilar"]:
        o.append(["Uyarı", u])
        o.cell(o.max_row, 1).fill = PatternFill("solid", fgColor="FFF4CE")
    o.append(["Not", "Projeksiyon basitleştirilmiştir: faiz giderinin vergi etkisi, yeni yatırımın getirisi, temettü ve ortak hareketleri "
                     "dikkate alınmaz. Rotatif krediler yalnız faiz servisiyle (anapara yenilenir varsayımıyla) hesaplanır. Kredi kararı ve "
                     "iç derecelendirme bankanın yetkili organlarına aittir."])
    o.cell(o.max_row, 2).alignment = UST
    o.merge_cells(start_row=o.max_row, start_column=2, end_row=o.max_row, end_column=3)

    p = wb.create_sheet("Projeksiyon (Baz)")
    alanlar = [("Net satışlar", "satis"), ("FAVÖK", "favok"), ("Vergi", "vergi"), ("Yatırım harcaması", "yatirim"), ("NİS artışı", "nis_degisim"),
               ("Borç servisine kullanılabilir nakit (CFADS)", "cfads"), ("Anapara", "anapara"), ("Faiz", "faiz"), ("Borç servisi", "servis"),
               ("DSCR", "dscr"), ("Yıl sonu finansal borç", "borc"), ("Net finansal borç / FAVÖK", "netborc_favok")]
    _baslik(p, ["Kalem"] + [f"{x['yil']}. yıl" for x in baz.projeksiyon], (42,) + (16,) * len(baz.projeksiyon))
    for ad, a in alanlar:
        p.append([ad] + [None if x[a] is None else (round(x[a], 2) if a in ("dscr", "netborc_favok") else round(x[a])) for x in baz.projeksiyon])
        for c in range(2, len(baz.projeksiyon) + 2):
            p.cell(p.max_row, c).number_format = "0.00" if a in ("dscr", "netborc_favok") else TL_BICIM
        if a in ("cfads", "dscr"):
            for h in p[p.max_row]:
                h.font = Font(bold=True)
        if a == "dscr":
            for c, x in enumerate(baz.projeksiyon, 2):
                p.cell(p.max_row, c).fill = PatternFill("solid", fgColor=RENK[sonuc_etiketi(x["dscr"], esik)])
    g = BarChart()
    g.title, g.height, g.width = "CFADS ve borç servisi", 8, 16
    g.add_data(Reference(p, min_col=1, max_col=1 + len(baz.projeksiyon), min_row=7), from_rows=True, titles_from_data=True)
    g.add_data(Reference(p, min_col=1, max_col=1 + len(baz.projeksiyon), min_row=10), from_rows=True, titles_from_data=True)
    g.set_categories(Reference(p, min_col=2, max_col=1 + len(baz.projeksiyon), min_row=1))
    p.add_chart(g, "A16")

    sn = wb.create_sheet("Senaryolar")
    _baslik(sn, ["Senaryo", "Satış %", "Marj puan", "Değişken faiz puan", "Kur %"] + [f"DSCR {i}. yıl" for i in range(1, s["yil"] + 1)]
            + ["En düşük", "Sonuç"], (24, 9, 10, 12, 8) + (11,) * s["yil"] + (10, 14))
    for ad, x in s["senaryolar"].items():
        sn.append([ad, x.tanim["satis"], x.tanim["marj"], x.tanim["faiz"], x.tanim["kur"]] +
                  [None if y["dscr"] is None else round(y["dscr"], 2) for y in x.projeksiyon] +
                  [None if min_dscr(x.projeksiyon) is None else round(min_dscr(x.projeksiyon), 2), x.sonuc])
        sn.cell(sn.max_row, sn.max_column).fill = PatternFill("solid", fgColor=RENK[x.sonuc])
    lc = LineChart()
    lc.title, lc.height, lc.width = "DSCR — senaryolar", 8, 18
    for i in range(len(s["senaryolar"])):
        lc.add_data(Reference(sn, min_col=6, max_col=5 + s["yil"], min_row=2 + i), from_rows=True, titles_from_data=False)
    lc.set_categories(Reference(sn, min_col=6, max_col=5 + s["yil"], min_row=1))
    for seri, ad in zip(lc.series, s["senaryolar"]):
        seri.tx = SeriesLabel(v=ad)
    sn.add_chart(lc, f"A{len(s['senaryolar']) + 4}")

    bs = wb.create_sheet("Borç Servisi")
    _baslik(bs, ["Kredi", "Para", "Bakiye", "Faiz %", "Kalan Vade", "Ödeme", "Faiz Tipi"] +
            [f"{i}. yıl {t}" for i in range(1, s["yil"] + 1) for t in ("anapara", "faiz")] + ["Not"],
            (34, 6, 14, 8, 10, 13, 10) + (13,) * (2 * s["yil"]) + (50,))
    for k in s["krediler"]:
        yillar = s["servis"].get(k.ad)
        if yillar is None:
            continue
        bs.append([k.ad, k.para, k.bakiye, k.faiz, k.vade, k.tip, "Değişken" if k.degisken else "Sabit"] +
                  [round(v_) for y in yillar for v_ in y[:2]] + ["; ".join(k.notlar)])
        for c in range(8, 8 + 2 * s["yil"]):
            bs.cell(bs.max_row, c).number_format = TL_BICIM
        bs.cell(bs.max_row, 3).number_format = TL_BICIM

    op = wb.create_sheet("Ödeme Planı")
    _baslik(op, ["Ay", "Anapara", "Faiz", "Taksit", "Kalan Bakiye"], (6, 15, 15, 15, 16))
    for i, (a, f, b) in enumerate(odeme_plani(oneri.bakiye, oneri.faiz, oneri.vade, oneri.tip, oneri.odemesiz), 1):
        op.append([i, round(a, 2), round(f, 2), round(a + f, 2), round(b, 2)])
        for c in range(2, 6):
            op.cell(op.max_row, c).number_format = "#,##0.00"

    va = wb.create_sheet("Varsayımlar")
    _baslik(va, ["Varsayım", "Değer", "Kaynak"], (34, 14, 50))
    for a, d, k in [("Satış büyümesi (yıllık, nominal)", v["buyume"], v["buyume_kaynak"]), ("Geçmiş bileşik satış büyümesi", v["cagr"], "mali tablolar"),
                    ("FAVÖK marjı", v["marj"], f"son yıl ({s['yillar'][-1]})"), ("Vergi / FAVÖK", v["vergi_orani"], "geçmiş yılların toplamı"),
                    ("Yatırım / satış", v["yatirim_orani"], "geçmiş yılların toplamı"), ("NİS / satış", v["nis_orani"], "son yıl"),
                    ("Hazır değerler (sabit)", v["hazir"], "son yıl")]:
        va.append([a, d, k])
        va.cell(va.max_row, 2).number_format = TL_BICIM if "Hazır" in a else "0.0%"
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


@dataclass
class Senaryo:
    tanim: dict
    projeksiyon: list[dict]
    sonuc: str


def calistir(mali: Path, borclar: Path | None, oneri_yolu: Path, cikti: Path, esik: float = 1.25, buyume: float | None = None,
             kurlar: dict[str, float] | None = None, senaryolar: list[dict] | None = None) -> dict:
    yillar, veri = mali_oku(mali)
    v = varsayimlar(yillar, veri, buyume)
    mevcut = borclari_oku(borclar) if borclar else []
    oneri, oneri_kur = oneri_oku(oneri_yolu)
    kur = {**oneri_kur, **(kurlar or {})}
    krediler = mevcut + [oneri]
    yil = max(1, math.ceil(max(k.vade for k in krediler) / 12))
    servis, uyarilar = yillik_servis(krediler, yil, kur)
    sonuc = {}
    for t in senaryolar or VARSAYILAN_SENARYOLAR:
        p = senaryo_calistir(v, krediler, kur, yil, t)
        sonuc[t["ad"]] = Senaryo(t, p, sonuc_etiketi(min_dscr(p), esik))
    if v["marj"] <= 0:
        uyarilar.append("Son yıl FAVÖK negatif veya sıfır: projeksiyon anlamlı değil")
    if any(k.tip == "rotatif" for k in mevcut):
        uyarilar.append("Rotatif krediler yalnız faiz servisiyle hesaplandı; yenilenmezse anapara da ödenecektir")
    s = {"yillar": yillar, "veri": veri, "varsayim": v, "krediler": krediler, "oneri": oneri, "servis": servis, "yil": yil, "esik": esik,
         "senaryolar": sonuc, "uyarilar": uyarilar, "azami": azami_kredi(v, mevcut, oneri, kur, yil, esik),
         "kirilma": kirilma_satis(v, krediler, kur, yil)}
    rapor_yaz(cikti, s)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Nakit akışı, borç servisi ve önerilen kredi yapısından DSCR ve senaryo analizi yapar.")
    p.add_argument("--mali", type=Path, default=ORNEK / "mali_tablolar.csv", help="Mali tablolar: Kalem × yıl (Net Satışlar, FAVÖK, Ödenen Vergi...)")
    p.add_argument("--borclar", type=Path, default=None, help="Mevcut krediler: Kredi, Bakiye, Para Birimi, Yıllık Faiz, Kalan Vade, Ödeme Tipi, Faiz Tipi")
    p.add_argument("--oneri", type=Path, default=ORNEK / "kredi_onerisi.json", help="Önerilen kredi yapısı (JSON)")
    p.add_argument("--min-dscr", default="1,25", help="Asgari DSCR eşiği (varsayılan 1,25 — örnektir)")
    p.add_argument("--buyume", help="Yıllık nominal satış büyümesi %% (varsayılan: geçmiş bileşik büyüme)")
    p.add_argument("--kur", nargs="*", default=[], metavar="PB=KUR", help="Döviz kurları, ör. EUR=48,05")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "geri_odeme_kapasitesi.xlsx")
    a = p.parse_args(argv)
    borclar = a.borclar or (ORNEK / "mevcut_borclar.csv" if a.mali == ORNEK / "mali_tablolar.csv" else None)
    kurlar = {}
    for x in a.kur:
        ad, _, d = x.partition("=")
        if sayi(d) is None:
            print(f"[X] --kur 'EUR=48,05' biçiminde olmalı: {x}")
            return 2
        kurlar[ad.strip().upper()] = sayi(d)
    s = calistir(a.mali, borclar, a.oneri, a.cikti, sayi(a.min_dscr) or 1.25, sayi(a.buyume) if a.buyume else None, kurlar)
    for ad, x in s["senaryolar"].items():
        m = min_dscr(x.projeksiyon)
        print(f"     {ad:<24} en düşük DSCR {'-' if m is None else f'{m:.2f}'.replace('.', ',')} · {x.sonuc}")
    print(f"[OK] Baz senaryoda taşınabilir azami kredi: {tl(s['azami'])} {s['oneri'].para} · "
          f"DSCR'yi 1,0'a indiren satış düşüşü: {'yok' if s['kirilma'] is None else '%' + str(s['kirilma']).replace('.', ',')}")
    for u in s["uyarilar"]:
        print(f"[!] {u}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
