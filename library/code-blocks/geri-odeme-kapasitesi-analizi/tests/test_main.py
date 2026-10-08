import sys
import tempfile
import unittest
from pathlib import Path

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))

import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"


class PlanTesti(unittest.TestCase):
    def test_esit_taksit(self):
        p = main.odeme_plani(100000, 12, 12, "esit-taksit")
        self.assertAlmostEqual(p[0][0] + p[0][1], 8884.88, places=2)          # %1 aylık, 12 ay annüite
        self.assertAlmostEqual(sum(a for a, _, _ in p), 100000, places=6)
        self.assertAlmostEqual(p[-1][2], 0, places=6)

    def test_odemesiz_ve_diger_tipler(self):
        p = main.odeme_plani(120000, 12, 12, "esit-anapara", odemesiz=6)
        self.assertEqual([round(a) for a, _, _ in p[:7]], [0] * 6 + [20000])
        self.assertAlmostEqual(p[0][1], 1200)                                  # ödemesiz dönemde faiz ödenir
        v = main.odeme_plani(50000, 24, 6, "vade-sonu")
        self.assertEqual([round(a) for a, _, _ in v], [0, 0, 0, 0, 0, 50000])
        r = main.odeme_plani(50000, 24, 6, "rotatif")
        self.assertEqual((sum(a for a, _, _ in r), round(r[0][1])), (0, 1000))

    def test_tip_coz(self):
        self.assertEqual([main.tip_coz(x) for x in ("Eşit taksit", "Eşit anapara", "Vade sonu", "Rotatif", "Spot", "")],
                         ["esit-taksit", "esit-anapara", "vade-sonu", "rotatif", "rotatif", "esit-taksit"])


class OrnekTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        with tempfile.TemporaryDirectory() as tmp:
            cls.s = main.calistir(ORNEK / "mali_tablolar.csv", ORNEK / "mevcut_borclar.csv", ORNEK / "kredi_onerisi.json", Path(tmp) / "g.xlsx")
            cls.wb = load_workbook(Path(tmp) / "g.xlsx")

    def test_varsayimlar(self):
        v = self.s["varsayim"]
        self.assertAlmostEqual(v["marj"], 30 / 132)
        self.assertAlmostEqual(v["buyume"], (132 / 80) ** 0.5 - 1)
        self.assertAlmostEqual(v["vergi_orani"], 10.2 / 70.5)
        self.assertAlmostEqual(v["yatirim_orani"], 7.5 / 317)
        self.assertAlmostEqual(v["nis_orani"], 11 / 132)

    def test_ilk_yil_elle(self):
        v = self.s["varsayim"]
        satis = 132e6 * (1 + v["buyume"])
        favok = satis * 30 / 132
        cfads = favok - favok * 10.2 / 70.5 - satis * 7.5 / 317 - (satis * 11 / 132 - 11e6)
        y1 = self.s["senaryolar"]["Baz"].projeksiyon[0]
        self.assertAlmostEqual(y1["cfads"], cfads, places=2)
        servis = sum(sum(a + f for a, f, _ in main.odeme_plani(k.bakiye, k.faiz, k.vade, k.tip, k.odemesiz)[:12])
                     * (48.05 if k.para == "EUR" else 1) for k in self.s["krediler"])
        self.assertAlmostEqual(y1["servis"], servis, places=2)
        self.assertAlmostEqual(y1["dscr"], cfads / servis)

    def test_senaryolar(self):
        sn = {a: (round(main.min_dscr(x.projeksiyon), 2), x.sonuc) for a, x in self.s["senaryolar"].items()}
        self.assertEqual(sn["Baz"], (1.14, "Sınırda"))
        self.assertEqual(sn["Kur +%30"][1], "Karşılamıyor")
        self.assertLess(sn["Birleşik stres"][0], sn["Baz"][0])
        self.assertLess(sn["Değişken faiz +10 puan"][0], sn["Baz"][0])

    def test_azami_ve_kirilma(self):
        self.assertTrue(10e6 < self.s["azami"] < 20e6)
        k = main.Kredi("x", self.s["azami"], "TL", 42, 36, "esit-taksit", False, 6, True)
        mevcut = [x for x in self.s["krediler"] if not x.yeni]
        m = main.min_dscr(main.senaryo_calistir(self.s["varsayim"], mevcut + [k], {"EUR": 48.05}, 3, main.VARSAYILAN_SENARYOLAR[0]))
        self.assertGreaterEqual(m, 1.25)
        self.assertAlmostEqual(self.s["kirilma"], 21.7, delta=0.1)

    def test_excel(self):
        self.assertEqual(self.wb.sheetnames, ["Özet", "Projeksiyon (Baz)", "Senaryolar", "Borç Servisi", "Ödeme Planı", "Varsayımlar"])
        self.assertEqual(self.wb["Ödeme Planı"].max_row, 37)


if __name__ == "__main__":
    unittest.main()
