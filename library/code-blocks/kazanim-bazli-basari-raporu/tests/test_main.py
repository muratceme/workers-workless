import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def q(soru, cevap, kazanim):
    return {"soru": soru, "cevap": main.katla(cevap), "cevap_ham": cevap, "kazanim": kazanim, "iptal": main.katla(cevap) in main.IPTAL}


def o(ad, sinif, *cevaplar):
    return {"ogrenci": ad, "sinif": sinif, "cevaplar": {str(i): c.lower() for i, c in enumerate(cevaplar, 1)}, "satir": 0}


SORU = lambda: {"1": q("1", "A", "K1"), "2": q("2", "B", "K1"), "3": q("3", "C", "K2")}   # noqa: E731
OGR = lambda: [o("Ali", "9A", "A", "B", "C"), o("Ayşe", "9A", "A", "C", ""), o("Can", "9B", "B", "B", "C")]   # noqa: E731


class AnalizTesti(unittest.TestCase):
    def test_kazanim_ve_sinif(self):
        r = main.analiz(OGR(), SORU())
        self.assertEqual([(k["kazanim"], k["soru"], k["sorular"], k["dogru"], k["yanlis"], k["bos"], k["basari"], k["sinif"], k["destek_ogrenci"], k["durum"]) for k in r["kazanimlar"]],
                         [("K1", 2, "1, 2", 4, 2, 0, D("66.67"), {"9A": D("75.00"), "9B": D("50.00")}, 0, "Yeterli"), ("K2", 1, "3", 2, 0, 1, D("66.67"), {"9A": D("50.00"), "9B": D("100.00")}, 1, "Yeterli")])
        self.assertEqual((r["genel"], r["siniflar"]), ({"ogrenci": 3, "soru": 3, "kazanim": 2, "basari": D("66.67"), "destek": 0}, [{"sinif": "9A", "ogrenci": 2, "basari": D("66.67")}, {"sinif": "9B", "ogrenci": 1, "basari": D("66.67")}]))
        self.assertEqual([(u["tur"], u["kim"]) for u in r["uyarilar"]], [("Tek soruyla ölçülen kazanım", "K2")])

    def test_ogrenci_ve_soru(self):
        r = main.analiz(OGR(), SORU())
        self.assertEqual([(x["ogrenci"], x["dogru"], x["yanlis"], x["bos"], x["basari"], x["ky"], x["destek"]) for x in r["ogrenciler"]],
                         [("Ali", 3, 0, 0, D("100.00"), {"K1": D("100.00"), "K2": D("100.00")}, []), ("Ayşe", 1, 1, 1, D("33.33"), {"K1": D("50.00"), "K2": D("0.00")}, ["K2"]),
                          ("Can", 2, 1, 0, D("66.67"), {"K1": D("50.00"), "K2": D("100.00")}, [])])
        self.assertEqual([(s["soru"], s["cevap"], s["dogru"], s["yanlis"], s["bos"], s["dogru_y"], s["bos_y"], s["celdirici"]) for s in r["sorular"]],
                         [("1", "A", 2, 1, 0, D("66.67"), D("0.00"), "B (1)"), ("2", "B", 2, 1, 0, D("66.67"), D("0.00"), "C (1)"), ("3", "C", 2, 0, 1, D("66.67"), D("33.33"), "")])

    def test_esik_iptal_ve_hatalar(self):
        r = main.analiz(OGR(), SORU(), D(70))
        self.assertEqual(([k["durum"] for k in r["kazanimlar"]], r["genel"]["destek"], [x["destek"] for x in r["ogrenciler"]], sorted(u["tur"] for u in r["uyarilar"])),
                         (["Destek gerekli", "Destek gerekli"], 2, [[], ["K1", "K2"], ["K1"]], ["Kazanım eşiğin altında", "Kazanım eşiğin altında", "Tek soruyla ölçülen kazanım"]))
        r = main.analiz(OGR(), SORU(), D(60))
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]), [("Sınıf düzeyinde düşük kazanım", "K1"), ("Sınıf düzeyinde düşük kazanım", "K2"), ("Tek soruyla ölçülen kazanım", "K2")])
        s = SORU() | {"2": q("2", "İptal", "K1"), "4": q("4", "A", "K9")}
        r = main.analiz(OGR(), s)
        self.assertEqual(([k["kazanim"] for k in r["kazanimlar"]], r["genel"]["soru"], r["kazanimlar"][0]["basari"], sorted(u["tur"] for u in r["uyarilar"])), (["K1", "K2"], 2, D("66.67"), ["Sınıf düzeyinde düşük kazanım"] + ["Tek soruyla ölçülen kazanım"] * 2 + ["İptal edilen soru"]))
        for og, so in (([], SORU()), (OGR(), {"2": q("2", "X", "K1")})):
            with self.assertRaises(ValueError):
                main.analiz(og, so)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "k.xlsx"
            r = main.calistir(ORNEK / "cevaplar.csv", ORNEK / "cevap_anahtari.csv", cikti)
            self.assertEqual([(k["kazanim"][:2], k["soru"], k["dogru"], k["basari"], k["sinif"], k["durum"]) for k in r["kazanimlar"]],
                             [("K1", 3, 19, D("79.17"), {"9-A": D("75.00"), "9-B": D("83.33")}, "Yeterli"), ("K2", 2, 9, D("56.25"), {"9-A": D("75.00"), "9-B": D("37.50")}, "Yeterli"),
                              ("K3", 2, 6, D("37.50"), {"9-A": D("50.00"), "9-B": D("25.00")}, "Destek gerekli")])
            self.assertEqual(([x["dogru"] for x in r["ogrenciler"]], r["genel"]["basari"], r["genel"]["soru"], [(u["onem"], u["tur"]) for u in r["uyarilar"]]),
                             ([7, 5, 4, 3, 3, 4, 2, 6], D("60.71"), 7, [("Yüksek", "Kazanım eşiğin altında"), ("Orta", "Sınıf düzeyinde düşük kazanım"), ("Bilgi", "İptal edilen soru")]))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Kazanım Özeti", "Öğrenci × Kazanım", "Soru Analizi", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            cy, ay = Path(tmp) / "c.csv", Path(tmp) / "a.csv"
            ay.write_text("Soru No;Cevap;Konu\nS1;a;K1\nSoru 2;B;K1\n2;C;K1\n3;;K2\n4;D;K2\n", encoding="utf-8")
            cy.write_text("Sınav sonuçları\nAd Soyad;Şube;Soru 1;S2;Not\nAli;A;a;B;x\n;;;;\nAyşe;;B;;\n", encoding="utf-8")
            sorular, uy = main.oku_anahtar(ay)
            self.assertEqual((list(sorular), sorular["1"]["cevap"], [u["tur"] for u in uy]), (["1", "2", "4"], "a", ["Anahtar satırı okunamadı"] * 2))
            ogr, uy = main.oku_cevaplar(cy, sorular)
            self.assertEqual(([(x["ogrenci"], x["sinif"], x["cevaplar"]) for x in ogr], [u["tur"] for u in uy]), ([("Ali", "A", {"1": "a", "2": "b"}), ("Ayşe", "Belirtilmemiş", {"1": "b", "2": ""})], ["Cevap sütunu olmayan soru"]))
            self.assertEqual(main.analiz(ogr, sorular)["genel"], {"ogrenci": 2, "soru": 2, "kazanim": 1, "basari": D("50.00"), "destek": 0})
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--cevaplar", str(cy), "--anahtar", str(ay), "--cikti", c]
            self.assertEqual(main.main(arg + ["--esik", "60"]), 0)
            self.assertEqual(main.main(arg + ["--esik", "0"]), 2)
            self.assertEqual(main.main(["--cevaplar", str(cy), "--cikti", c]), 2)
            self.assertEqual(main.main(["--cevaplar", str(cy), "--anahtar", str(Path(tmp) / "yok.csv")]), 1)
            cy.write_text("Ad;Sınıf;1\nAli;A;a\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            cy.write_text("Öğrenci;Sınıf;Q9\nAli;A;a\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
