"""
Kazanım Bazlı Başarı Raporu — Gerisi Bende kod bloğu
Ölçme Değerlendirme ve Rehberlik › Ölçme Değerlendirme Uzmanı

Çoktan seçmeli sınav cevaplarını cevap anahtarı ve soru-kazanım eşleşmesiyle değerlendirir; sınıf ve öğrenci bazında kazanım başarı yüzdelerini raporlar.
  Doğru / yanlış / boş = öğrencinin cevabı anahtardaki doğru cevapla aynıysa doğru; boşsa boş; değilse yanlış          (büyük-küçük harf ayrılmaz)
  Kazanım başarısı %   = kazanımın sorularındaki doğru sayısı ÷ (öğrenci sayısı × soru sayısı) × 100
  Sınıf × kazanım      = aynı oran, yalnız o sınıfın öğrencileriyle
  Öğrenci × kazanım    = öğrencinin o kazanımdaki doğru sayısı ÷ kazanımın soru sayısı × 100
  Soru analizi         = doğru %, boş %, en çok işaretlenen yanlış seçenek
  Destek gerekli       = başarısı --esik (%50) altında kalan kazanım (genel, sınıf ve öğrenci düzeyinde)
  İptal soru           = anahtarda doğru cevabı "İptal" veya "X" olan soru değerlendirmeye alınmaz
Yanlışın doğruyu götürdüğü net hesabı ve puanlama yapılmaz; rapor kazanım düzeyinde doğru oranına bakar. Az soruyla ölçülen kazanımda yüzde kaba bir göstergedir.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


ANAHTAR = {"soru": ("soru", "soru no", "no"), "cevap": ("dogru cevap", "cevap", "anahtar"), "kazanim": ("kazanim", "kazanim kodu", "konu")}
OGRENCI, SINIF = ("ogrenci", "ogrenci adi", "ad soyad", "adi soyadi", "ogrenci no"), ("sinif", "sube", "sinif sube")
IPTAL = {"iptal", "x"}


def soru_anahtari(x) -> str:
    if isinstance(x, float) and x == int(x):
        x = int(x)
    s = katla(x)
    r = re.findall(r"\d+", s)
    return str(int(r[-1])) if r else s


def oku_anahtar(yol: Path) -> tuple[dict, list[dict]]:
    uy, sorular = [], {}
    for r in kayitlar(yol, ANAHTAR, ("soru", "cevap", "kazanim")):
        s = soru_anahtari(r["soru"])
        if not s:
            continue
        c, k = katla(r["cevap"]), metin(r["kazanim"])
        if not c or not k or s in sorular:
            uy.append(uyari("Yüksek", "Anahtar satırı okunamadı", f"Soru {metin(r['soru'])} (satır {r['_satir']})", "Doğru cevap veya kazanım boş ya da soru numarası tekrar ediyor. Soru değerlendirmeye alınmadı."))
            continue
        sorular[s] = {"soru": s, "cevap": c, "cevap_ham": metin(r["cevap"]), "kazanim": k, "iptal": c in IPTAL}
    return sorular, uy


def oku_cevaplar(yol: Path, sorular: dict) -> tuple[list[dict], list[dict]]:
    uy, satirlar = [], tablo_oku(yol)
    bi = next((i for i, r in enumerate(satirlar[:15]) if any(katla(c) in OGRENCI for c in r)), None)
    if bi is None:
        raise ValueError(f"{yol.name}: başlık satırı bulunamadı; 'Öğrenci' sütunu gerekli.")
    b = [katla(c) for c in satirlar[bi]]
    oj = next(j for j, c in enumerate(b) if c in OGRENCI)
    sj = next((j for j, c in enumerate(b) if c in SINIF), None)
    sutun = {}
    for j, c in enumerate(satirlar[bi]):
        if j not in (oj, sj) and c not in (None, "") and soru_anahtari(c) in sorular:
            sutun.setdefault(soru_anahtari(c), j)
    eksik = [s for s in sorular if s not in sutun]
    if len(eksik) == len(sorular):
        raise ValueError(f"{yol.name}: cevap dosyasında anahtardaki soru numaralarıyla eşleşen sütun yok.")
    if eksik:
        uy.append(uyari("Yüksek", "Cevap sütunu olmayan soru", "Cevap dosyası", f"Anahtarda olup cevap dosyasında sütunu bulunmayan sorular: {', '.join(eksik[:15])}. Bu sorular değerlendirilmedi."))
    ogrenciler = []
    for n, r in enumerate(satirlar[bi + 1:], bi + 2):
        ad = metin(r[oj]) if oj < len(r) else ""
        if isinstance(r[oj] if oj < len(r) else None, float) and r[oj] == int(r[oj]):
            ad = str(int(r[oj]))
        if not ad:
            continue
        ogrenciler.append({"ogrenci": ad, "sinif": (metin(r[sj]) if sj is not None and sj < len(r) else "") or "Belirtilmemiş", "cevaplar": {s: katla(r[j]) if j < len(r) else "" for s, j in sutun.items()}, "satir": n})
    return ogrenciler, uy


def analiz(ogrenciler: list[dict], sorular: dict, esik: Decimal = Decimal(50)) -> dict:
    if not ogrenciler:
        raise ValueError("Değerlendirilecek öğrenci yok.")
    uy = []
    iptal = [s for s, q in sorular.items() if q["iptal"]]
    if iptal:
        uy.append(uyari("Bilgi", "İptal edilen soru", "Anahtar", f"Değerlendirmeye alınmayan sorular: {', '.join(iptal)}."))
    gecerli = [q for s, q in sorular.items() if not q["iptal"] and any(s in o["cevaplar"] for o in ogrenciler)]
    if not gecerli:
        raise ValueError("Değerlendirilecek soru yok.")
    kazanimlar = list(dict.fromkeys(q["kazanim"] for q in gecerli))
    k_soru = {k: [q["soru"] for q in gecerli if q["kazanim"] == k] for k in kazanimlar}
    yuzde = lambda a, b: yuvarla(Decimal(a) / b * 100) if b else None   # noqa: E731
    ogr = []
    for o in ogrenciler:
        d = {q["soru"]: ("B" if not o["cevaplar"].get(q["soru"]) else "D" if o["cevaplar"][q["soru"]] == q["cevap"] else "Y") for q in gecerli}
        kd = {k: sum(1 for s in k_soru[k] if d[s] == "D") for k in kazanimlar}
        ky = {k: yuzde(kd[k], len(k_soru[k])) for k in kazanimlar}
        ogr.append({"ogrenci": o["ogrenci"], "sinif": o["sinif"], "sonuc": d, "dogru": sum(kd.values()), "yanlis": sum(1 for v in d.values() if v == "Y"), "bos": sum(1 for v in d.values() if v == "B"), "kd": kd, "ky": ky,
                    "basari": yuzde(sum(kd.values()), len(gecerli)), "destek": [k for k in kazanimlar if ky[k] < esik]})
    siniflar = sorted({o["sinif"] for o in ogr})

    def kazanim_orani(grup: list[dict], k: str) -> Decimal | None:
        return yuzde(sum(o["kd"][k] for o in grup), len(grup) * len(k_soru[k]))

    kaz = []
    for k in kazanimlar:
        n = len(ogr) * len(k_soru[k])
        dogru, bos = sum(o["kd"][k] for o in ogr), sum(1 for o in ogr for s in k_soru[k] if o["sonuc"][s] == "B")
        x = {"kazanim": k, "soru": len(k_soru[k]), "sorular": ", ".join(k_soru[k]), "dogru": dogru, "yanlis": n - dogru - bos, "bos": bos, "basari": yuzde(dogru, n), "sinif": {s: kazanim_orani([o for o in ogr if o["sinif"] == s], k) for s in siniflar},
             "destek_ogrenci": sum(1 for o in ogr if k in o["destek"])}
        x["durum"] = "Destek gerekli" if x["basari"] < esik else "Yeterli"
        kaz.append(x)
        if x["durum"] == "Destek gerekli":
            uy.append(uyari("Yüksek", "Kazanım eşiğin altında", k, f"Genel başarı %{x['basari']} (eşik %{esik.normalize():f}); {x['destek_ogrenci']} öğrenci eşiğin altında. Sorular: {x['sorular']}."))
        else:
            dusuk = [s for s in siniflar if x["sinif"][s] < esik]
            if dusuk:
                ayrinti = ", ".join(f"{s} (%{x['sinif'][s]})" for s in dusuk)
                uy.append(uyari("Orta", "Sınıf düzeyinde düşük kazanım", k, f"Genel başarı %{x['basari']} ama şu sınıflarda eşiğin altında: {ayrinti}."))
        if len(k_soru[k]) == 1:
            uy.append(uyari("Bilgi", "Tek soruyla ölçülen kazanım", k, "Kazanım yalnız bir soruyla ölçülmüş; yüzde kaba bir göstergedir."))
    soru_analiz = []
    for q in gecerli:
        s, n = q["soru"], len(ogr)
        yanlislar = defaultdict(int)
        for o, ham in zip(ogr, ogrenciler):
            if o["sonuc"][s] == "Y":
                yanlislar[ham["cevaplar"][s].upper()] += 1
        d, b = sum(1 for o in ogr if o["sonuc"][s] == "D"), sum(1 for o in ogr if o["sonuc"][s] == "B")
        celdirici = min(yanlislar.items(), key=lambda t: (-t[1], t[0])) if yanlislar else None
        soru_analiz.append({"soru": s, "kazanim": q["kazanim"], "cevap": q["cevap_ham"].upper(), "dogru": d, "yanlis": n - d - b, "bos": b, "dogru_y": yuzde(d, n), "bos_y": yuzde(b, n), "celdirici": f"{celdirici[0]} ({celdirici[1]})" if celdirici else ""})
    sinif_ozet = [{"sinif": s, "ogrenci": len(g), "basari": yuzde(sum(o["dogru"] for o in g), len(g) * len(gecerli))} for s in siniflar for g in [[o for o in ogr if o["sinif"] == s]]]
    genel = {"ogrenci": len(ogr), "soru": len(gecerli), "kazanim": len(kazanimlar), "basari": yuzde(sum(o["dogru"] for o in ogr), len(ogr) * len(gecerli)), "destek": sum(1 for x in kaz if x["durum"] == "Destek gerekli")}
    return {"ogrenciler": ogr, "kazanimlar": kaz, "sorular": soru_analiz, "siniflar": sinif_ozet, "genel": genel, "esik": esik, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb, esik = Workbook(), r["esik"]
    siniflar = [s["sinif"] for s in r["siniflar"]]
    kaz = [k["kazanim"] for k in r["kazanimlar"]]
    renk = lambda v: KIRMIZI if v is not None and v < esik else None   # noqa: E731
    ws = wb.active
    ws.title = "Kazanım Özeti"
    _baslik(ws, ["Kazanım", "Soru Sayısı", "Sorular", "Doğru", "Yanlış", "Boş", "Genel Başarı %"] + [f"{s} %" for s in siniflar] + ["Eşik Altı Öğrenci", "Durum"], (40, 8, 16, 8, 8, 8, 11) + (10,) * len(siniflar) + (10, 16), "B2")
    for k in r["kazanimlar"]:
        _satir(ws, [k["kazanim"], k["soru"], k["sorular"], k["dogru"], k["yanlis"], k["bos"], k["basari"]] + [k["sinif"][s] for s in siniflar] + [k["destek_ogrenci"] or None, k["durum"]], para_s=tuple(range(7, 8 + len(siniflar))),
               dolgu={7: renk(k["basari"]), 9 + len(siniflar): KIRMIZI if k["durum"] == "Destek gerekli" else YESIL} | {8 + j: renk(k["sinif"][s]) for j, s in enumerate(siniflar)})
    n = _satir(ws, ["Genel", r["genel"]["soru"], "", None, None, None, r["genel"]["basari"]] + [s["basari"] for s in r["siniflar"]], para_s=tuple(range(7, 8 + len(siniflar))))
    for c in ws[n]:
        c.font = Font(bold=True)
    ws.append([])
    ws.append([f"Başarı % = doğru sayısı ÷ (öğrenci × soru). Kırmızı hücreler %{esik.normalize():f} eşiğinin altındadır. Net veya puan hesabı yapılmaz."])
    ws = wb.create_sheet("Öğrenci × Kazanım")
    _baslik(ws, ["Öğrenci", "Sınıf", "Doğru", "Yanlış", "Boş", "Başarı %"] + kaz + ["Destek Gereken Kazanımlar"], (28, 10, 8, 8, 8, 10) + (14,) * len(kaz) + (50,), "C2")
    for o in sorted(r["ogrenciler"], key=lambda o: (o["sinif"], o["ogrenci"])):
        _satir(ws, [o["ogrenci"], o["sinif"], o["dogru"], o["yanlis"], o["bos"], o["basari"]] + [o["ky"][k] for k in kaz] + [", ".join(o["destek"])], para_s=tuple(range(6, 7 + len(kaz))), dolgu={7 + j: renk(o["ky"][k]) for j, k in enumerate(kaz)},
               sar=(7 + len(kaz),))
    ws = wb.create_sheet("Soru Analizi")
    _baslik(ws, ["Soru", "Kazanım", "Doğru Cevap", "Doğru", "Yanlış", "Boş", "Doğru %", "Boş %", "En Çok İşaretlenen Yanlış"], (8, 40, 10, 8, 8, 8, 10, 9, 16))
    for s in r["sorular"]:
        _satir(ws, [s["soru"], s["kazanim"], s["cevap"], s["dogru"], s["yanlis"], s["bos"], s["dogru_y"], s["bos_y"], s["celdirici"]], para_s=(7, 8), dolgu={7: renk(s["dogru_y"])})
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(cevap_yol: Path, anahtar_yol: Path, cikti: Path, esik: Decimal = Decimal(50)) -> dict:
    sorular, uy = oku_anahtar(anahtar_yol)
    if not sorular:
        raise ValueError("Cevap anahtarında geçerli soru yok.")
    ogrenciler, u2 = oku_cevaplar(cevap_yol, sorular)
    r = analiz(ogrenciler, sorular, esik)
    uy += u2 + r["uyarilar"]
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Sınav cevaplarından sınıf ve öğrenci bazında kazanım başarı raporu üretir.")
    p.add_argument("--cevaplar", type=Path, default=ORNEK / "cevaplar.csv", help="Öğrenci, Sınıf ve her soru için bir sütun (başlık soru numarası; hücre işaretlenen seçenek)")
    p.add_argument("--anahtar", type=Path, help="Soru, Doğru Cevap, Kazanım")
    p.add_argument("--esik", default="50", help="Bu %%'nin altındaki kazanımlar 'destek gerekli' (varsayılan 50)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kazanim_basari_raporu.xlsx")
    a = p.parse_args(argv)
    anahtar = a.anahtar or (ORNEK / "cevap_anahtari.csv" if a.cevaplar == ORNEK / "cevaplar.csv" else None)
    e = para(a.esik)
    if anahtar is None or e is None or not 0 < e <= 100:
        print("[X] --anahtar dosyası verilmeli; --esik 0–100 arasında olmalı")
        return 2
    for y in (a.cevaplar, anahtar):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.cevaplar, anahtar, a.cikti, e)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    g = r["genel"]
    print(f"[OK] {g['ogrenci']} öğrenci · {g['soru']} soru · {g['kazanim']} kazanım · genel başarı %{g['basari']} · {g['destek']} kazanım eşiğin altında")
    for k in r["kazanimlar"]:
        if k["durum"] == "Destek gerekli":
            print(f"[!] {k['kazanim']}: %{k['basari']} ({k['destek_ogrenci']} öğrenci eşiğin altında)")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
