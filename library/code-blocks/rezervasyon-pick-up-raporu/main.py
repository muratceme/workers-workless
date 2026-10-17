"""
Rezervasyon Pick-up Raporu — Gerisi Bende kod bloğu
Rezervasyon ve Gelir Yönetimi › Gelir Müdürü (Revenue Manager)

Gelecek konaklama tarihleri için iki rezervasyon dökümü arasındaki artışı (pick-up) hesaplar; geçen yıl ve bütçeyle karşılaştırır.
  Pick-up (oda)     = bugünkü dökümdeki satılan oda − önceki dökümdeki satılan oda                (negatifse iptal veya değişiklik ağır basmış)
  Pick-up (gelir)   = bugünkü gelir − önceki gelir
  ADR               = gelir ÷ satılan oda;   pick-up ADR = pick-up gelir ÷ pick-up oda              (yalnız pick-up oda > 0 ise)
  Doluluk %         = satılan oda ÷ kapasite × 100                                                (--kapasite verilirse)
  Geçen yıla fark   = satılan oda − geçen yıl satılan oda;   fark % = fark ÷ geçen yıl × 100
  Bütçe gerçekleşme = satılan oda ÷ bütçe oda × 100;   bütçeye kalan = bütçe oda − satılan oda
  Aylık özet        = ay içindeki tarihlerin toplamı; ADR ve oranlar toplamlardan hesaplanır
Dökümler aynı esasta olmalıdır (aynı oda tipleri, KDV ve kahvaltı dahil veya hariç aynı gelir tanımı). Geçen yıl verisini hangi güne eşleyeceğinize (aynı tarih veya aynı hafta günü) siz karar verirsiniz.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"tarih": ("konaklama tarihi", "tarih", "gun"), "oda": ("oda", "oda sayisi", "satilan oda", "rezerve oda", "oda gece"), "gelir": ("gelir", "oda geliri", "konaklama geliri", "ciro")}
Z = Decimal(0)
GUNLER = ("Pzt", "Sal", "Çar", "Per", "Cum", "Cmt", "Paz")
AYLAR = ("Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık")


def oku(yol: Path, ad: str) -> tuple[dict, list[dict]]:
    uy, d = [], {}
    for r in kayitlar(yol, SUTUN, ("tarih", "oda")):
        if not metin(r["tarih"]) and not metin(r["oda"]):
            continue
        t, o, g = tarih(r["tarih"]), para(r["oda"]), para(r["gelir"]) if metin(r["gelir"]) else Z
        if t is None or o is None or o < 0 or g is None or g < 0:
            uy.append(uyari("Yüksek", "Satır okunamadı", f"{ad} (satır {r['_satir']})", f"Tarih '{metin(r['tarih'])}', oda '{metin(r['oda'])}', gelir '{metin(r['gelir'])}'. Satır alınmadı."))
            continue
        x = d.setdefault(t, {"oda": Z, "gelir": Z})                                              # aynı tarih birden çok satırdaysa (oda tipi, segment) toplanır
        x["oda"] += o
        x["gelir"] += g
    return d, uy


def bol(a: Decimal | None, b: Decimal | None, carpan: int = 1) -> Decimal | None:
    return yuvarla(a / b * carpan) if a is not None and b else None


def analiz(bugun: dict, onceki: dict, gecen: dict | None = None, butce: dict | None = None, kapasite: int | None = None) -> dict:
    if not bugun:
        raise ValueError("Bugünkü dökümde geçerli satır yok.")
    if not onceki:
        raise ValueError("Önceki dökümde geçerli satır yok.")
    uy, gunler, aylar = [], [], {}
    alanlar = ("oda", "gelir", "o_oda", "o_gelir", "p_oda", "p_gelir", "ly_oda", "ly_gelir", "b_oda", "b_gelir", "oda_ly", "oda_b")

    def oranlar(x: dict, gun_sayisi: int = 1) -> dict:
        return {"adr": bol(x["gelir"], x["oda"]), "p_adr": bol(x["p_gelir"], x["p_oda"]) if x["p_oda"] > 0 else None, "doluluk": bol(x["oda"], Decimal(kapasite * gun_sayisi), 100) if kapasite else None,
                "ly_fark": x["oda_ly"] - x["ly_oda"] if x["ly_oda"] is not None else None, "ly_fark_y": bol(x["oda_ly"] - x["ly_oda"], x["ly_oda"], 100) if x["ly_oda"] else None,
                "b_gerc": bol(x["oda_b"], x["b_oda"], 100), "b_kalan": x["b_oda"] - x["oda_b"] if x["b_oda"] is not None else None}       # karşılaştırma yalnız geçen yıl / bütçe verisi olan günler üzerinden

    for t in sorted(set(bugun) | set(onceki)):
        b, o = bugun.get(t), onceki.get(t, {"oda": Z, "gelir": Z})
        if b is None:
            uy.append(uyari("Bilgi", "Tarih bugünkü dökümde yok", f"{t:%d.%m.%Y}", f"Önceki dökümde {o['oda']:f} oda vardı; bugünkü dökümde tarih bulunmadığı için 0 sayıldı."))
            b = {"oda": Z, "gelir": Z}
        ly, bt = (gecen or {}).get(t), (butce or {}).get(t)
        x = {"tarih": t, "gun": GUNLER[t.weekday()], "oda": b["oda"], "gelir": b["gelir"], "o_oda": o["oda"], "o_gelir": o["gelir"], "p_oda": b["oda"] - o["oda"], "p_gelir": b["gelir"] - o["gelir"],
             "ly_oda": ly["oda"] if ly else None, "ly_gelir": ly["gelir"] if ly else None, "b_oda": bt["oda"] if bt else None, "b_gelir": bt["gelir"] if bt else None,
             "oda_ly": b["oda"] if ly else None, "oda_b": b["oda"] if bt else None}
        x |= oranlar(x)
        isaret = []
        if x["p_oda"] < 0:
            isaret.append("Negatif pick-up")
            uy.append(uyari("Orta", "Negatif pick-up", f"{t:%d.%m.%Y}", f"Satılan oda {o['oda']:f} → {b['oda']:f} ({x['p_oda']:f}). İptal, grup düşüşü veya tarih değişikliği olabilir."))
        if kapasite and x["oda"] > kapasite:
            isaret.append("Kapasite üstü")
            uy.append(uyari("Yüksek", "Kapasite üstü satış", f"{t:%d.%m.%Y}", f"Satılan oda {b['oda']:f}, kapasite {kapasite}. Fazla rezervasyon (overbooking) bilinçli değilse kontrol edin."))
        gunler.append(x | {"isaret": isaret})
        a = aylar.setdefault((t.year, t.month), {"ay": f"{AYLAR[t.month - 1]} {t.year}", "gun_sayisi": 0} | {k: None for k in alanlar})
        a["gun_sayisi"] += 1
        for k in alanlar:
            if x[k] is not None:
                a[k] = (a[k] or Z) + x[k]
    toplam = {"ay": "Toplam", "gun_sayisi": len(gunler)}
    for k in alanlar:
        d = [x[k] for x in gunler if x[k] is not None]
        toplam[k] = sum(d, Z) if d else None
    ay_listesi = [a | oranlar(a, a["gun_sayisi"]) for _, a in sorted(aylar.items())]
    return {"gunler": gunler, "aylar": ay_listesi, "toplam": toplam | oranlar(toplam, toplam["gun_sayisi"]), "kapasite": kapasite, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Günlük Pick-up"
    ust = ["Satılan Oda", "Önceki Oda", "Pick-up Oda", "Gelir", "Pick-up Gelir", "ADR", "Pick-up ADR", "Doluluk %", "Geçen Yıl Oda", "Geçen Yıla Fark", "Fark %", "Bütçe Oda", "Bütçe Gerçekleşme %", "Bütçeye Kalan"]
    gen = (10, 10, 10, 15, 14, 11, 11, 9, 10, 10, 8, 10, 12, 10)

    def degerler(x: dict) -> list:
        return [x["oda"], x["o_oda"], x["p_oda"], x["gelir"], x["p_gelir"], x["adr"], x["p_adr"], x["doluluk"], x["ly_oda"], x["ly_fark"], x["ly_fark_y"], x["b_oda"], x["b_gerc"], x["b_kalan"]]

    _baslik(ws, ["Tarih", "Gün"] + ust + ["İşaret"], (12, 6) + gen + (22,), "C2")
    for x in r["gunler"]:
        _satir(ws, [x["tarih"], x["gun"]] + degerler(x) + [", ".join(x["isaret"])], para_s=(6, 7, 8, 9, 10, 13, 15), tarih_s=(1,),
               dolgu={5: KIRMIZI if x["p_oda"] < 0 else YESIL if x["p_oda"] > 0 else None, 17: SARI if x["isaret"] else None, 2: MAVI if x["gun"] in ("Cum", "Cmt") else None})
    ws.append([])
    ws.append(["Pick-up = bugünkü döküm − önceki döküm. Negatif pick-up iptal veya değişiklik demektir. Doluluk yalnız --kapasite verilirse hesaplanır."])
    ws = wb.create_sheet("Aylık Özet")
    _baslik(ws, ["Ay", "Gün Sayısı"] + ust, (16, 9) + gen)
    for a in r["aylar"] + [r["toplam"]]:
        n = _satir(ws, [a["ay"], a["gun_sayisi"]] + degerler(a), para_s=(6, 7, 8, 9, 10, 13, 15))
        if a is r["toplam"]:
            for c in ws[n]:
                c.font = Font(bold=True)
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(bugun_yol: Path, onceki_yol: Path, cikti: Path, gecen_yol: Path | None = None, butce_yol: Path | None = None, kapasite: int | None = None) -> dict:
    bugun, uy = oku(bugun_yol, "Bugünkü döküm")
    onceki, u2 = oku(onceki_yol, "Önceki döküm")
    uy += u2
    gecen = butce = None
    if gecen_yol:
        gecen, u3 = oku(gecen_yol, "Geçen yıl")
        uy += u3
    if butce_yol:
        butce, u4 = oku(butce_yol, "Bütçe")
        uy += u4
    r = analiz(bugun, onceki, gecen, butce, kapasite)
    uy += r["uyarilar"]
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="İki rezervasyon dökümü arasındaki pick-up'ı hesaplar; geçen yıl ve bütçeyle karşılaştırır.")
    p.add_argument("--bugun", type=Path, default=ORNEK / "bugun.csv", help="Bugünkü döküm: Konaklama Tarihi, Satılan Oda, Gelir")
    p.add_argument("--onceki", type=Path, help="Önceki döküm (ör. 7 gün önce): aynı sütunlar")
    p.add_argument("--gecen-yil", type=Path, help="İsteğe bağlı: geçen yılın aynı döneme ait verisi, bu yılın tarihleriyle")
    p.add_argument("--butce", type=Path, help="İsteğe bağlı: bütçe, aynı sütunlar")
    p.add_argument("--kapasite", type=int, help="Satılabilir oda sayısı (doluluk için)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "pick_up_raporu.xlsx")
    a = p.parse_args(argv)
    ornek = a.bugun == ORNEK / "bugun.csv"
    onceki, gecen, butce = a.onceki or (ORNEK / "onceki.csv" if ornek else None), a.gecen_yil or (ORNEK / "gecen_yil.csv" if ornek else None), a.butce or (ORNEK / "butce.csv" if ornek else None)
    kapasite = a.kapasite if a.kapasite is not None else 120 if ornek else None
    if onceki is None or (kapasite is not None and kapasite <= 0):
        print("[X] --onceki dosyası verilmeli; --kapasite pozitif tam sayı olmalı")
        return 2
    for y in (a.bugun, onceki, gecen, butce):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.bugun, onceki, a.cikti, gecen, butce, kapasite)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {t['gun_sayisi']} gün · satılan {t['oda']:,.0f} oda · pick-up {t['p_oda']:+,.0f} oda, {t['p_gelir']:+,.2f} gelir · ADR {t['adr'] or 0:,.2f}" + (f" · doluluk %{t['doluluk']}" if t["doluluk"] is not None else ""))
    for x in r["gunler"]:
        if x["isaret"]:
            print(f"[!] {x['tarih']:%d.%m.%Y} {x['gun']}: {', '.join(x['isaret'])} ({x['p_oda']:+f} oda)")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
