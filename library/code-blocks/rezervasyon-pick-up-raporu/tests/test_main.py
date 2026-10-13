import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
A, B, C, E = date(2026, 10, 30), date(2026, 10, 31), date(2026, 11, 1), date(2026, 11, 2)


def v(**kw):
    return {t: {"oda": D(o), "gelir": D(g)} for t, (o, g) in kw.items()}


BUGUN = lambda: {A: {"oda": D(10), "gelir": D(1000)}, B: {"oda": D(8), "gelir": D(900)}, C: {"oda": D(5), "gelir": D(500)}}   # noqa: E731
ONCEKI = lambda: {A: {"oda": D(7), "gelir": D(640)}, B: {"oda": D(9), "gelir": D(1000)}, E: {"oda": D(2), "gelir": D(200)}}   # noqa: E731
GECEN, BUTCE = {A: {"oda": D(8), "gelir": D(700)}}, {A: {"oda": D(12), "gelir": D(1300)}, B: {"oda": D(10), "gelir": D(1000)}}


class AnalizTesti(unittest.TestCase):
    def test_gunluk(self):
        r = main.analiz(BUGUN(), ONCEKI(), GECEN, BUTCE, 9)
        self.assertEqual([(x["tarih"], x["gun"], x["oda"], x["p_oda"], x["p_gelir"], x["adr"], x["p_adr"], x["doluluk"], x["isaret"]) for x in r["gunler"]],
                         [(A, "Cum", D(10), D(3), D(360), D("100.00"), D("120.00"), D("111.11"), ["Kapasite üstü"]), (B, "Cmt", D(8), D(-1), D(-100), D("112.50"), None, D("88.89"), ["Negatif pick-up"]),
                          (C, "Paz", D(5), D(5), D(500), D("100.00"), D("100.00"), D("55.56"), []), (E, "Pzt", D(0), D(-2), D(-200), None, None, D("0.00"), ["Negatif pick-up"])])
        self.assertEqual([(x["ly_oda"], x["ly_fark"], x["ly_fark_y"], x["b_oda"], x["b_gerc"], x["b_kalan"]) for x in r["gunler"]][:3],
                         [(D(8), D(2), D("25.00"), D(12), D("83.33"), D(2)), (None, None, None, D(10), D("80.00"), D(2)), (None, None, None, None, None, None)])
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]), [("Kapasite üstü satış", "30.10.2026"), ("Negatif pick-up", "02.11.2026"), ("Negatif pick-up", "31.10.2026"), ("Tarih bugünkü dökümde yok", "02.11.2026")])

    def test_aylik_ve_toplam(self):
        r = main.analiz(BUGUN(), ONCEKI(), GECEN, BUTCE, 9)
        al = ("ay", "gun_sayisi", "oda", "gelir", "o_oda", "p_oda", "p_gelir", "adr", "p_adr", "doluluk", "ly_oda", "ly_fark", "ly_fark_y", "b_oda", "b_gerc", "b_kalan")
        self.assertEqual([tuple(a[k] for k in al) for a in r["aylar"] + [r["toplam"]]],
                         [("Ekim 2026", 2, D(18), D(1900), D(16), D(2), D(260), D("105.56"), D("130.00"), D("100.00"), D(8), D(2), D("25.00"), D(22), D("81.82"), D(4)),
                          ("Kasım 2026", 2, D(5), D(500), D(2), D(3), D(300), D("100.00"), D("100.00"), D("27.78"), None, None, None, None, None, None),
                          ("Toplam", 4, D(23), D(2400), D(18), D(5), D(560), D("104.35"), D("112.00"), D("63.89"), D(8), D(2), D("25.00"), D(22), D("81.82"), D(4))])
        r = main.analiz(BUGUN(), ONCEKI())
        self.assertEqual((r["toplam"]["doluluk"], r["toplam"]["ly_fark"], r["toplam"]["b_gerc"], [x["isaret"] for x in r["gunler"]]), (None, None, None, [[], ["Negatif pick-up"], [], ["Negatif pick-up"]]))
        for a, b in (({}, ONCEKI()), (BUGUN(), {})):
            with self.assertRaises(ValueError):
                main.analiz(a, b)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "p.xlsx"
            r = main.calistir(ORNEK / "bugun.csv", ORNEK / "onceki.csv", cikti, ORNEK / "gecen_yil.csv", ORNEK / "butce.csv", 120)
            t = r["toplam"]
            self.assertEqual((t["gun_sayisi"], t["oda"], t["gelir"], t["p_oda"], t["p_gelir"], t["adr"], t["p_adr"], t["doluluk"], t["ly_fark"], t["ly_fark_y"], t["b_gerc"], t["b_kalan"]),
                             (10, D(758), D("2263800.00"), D(76), D("250800.00"), D("2986.54"), D("3300.00"), D("63.17"), D(46), D("6.46"), D("93.01"), D(57)))
            self.assertEqual(([(x["gun"], x["p_oda"], x["doluluk"]) for x in r["gunler"]][:5], [(u["tur"], u["kim"]) for u in r["uyarilar"]], [a["ay"] for a in r["aylar"]]),
                             ([("Sal", D(5), D("79.17")), ("Çar", D(8), D("73.33")), ("Per", D(-2), D("63.33")), ("Cum", D(14), D("91.67")), ("Cmt", D(18), D("98.33"))], [("Negatif pick-up", "22.10.2026")], ["Ekim 2026"]))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Günlük Pick-up", "Aylık Özet", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            b, o = Path(tmp) / "b.csv", Path(tmp) / "o.csv"
            b.write_text("Tarih;Oda;Gelir\n01.11.2026;10;12.500,50\n01.11.2026;5;\n02.11.2026;x;5\n03.11;4;4\n04.11.2026;-1;0\n", encoding="utf-8")
            o.write_text("Tarih;Oda Sayısı;Oda Geliri\n01.11.2026;12;10000\n", encoding="utf-8")
            d, uy = main.oku(b, "Bugünkü döküm")
            self.assertEqual((d, [(u["tur"], u["kim"]) for u in uy]), ({date(2026, 11, 1): {"oda": D(15), "gelir": D("12500.50")}}, [("Satır okunamadı", f"Bugünkü döküm (satır {n})") for n in (4, 5, 6)]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--bugun", str(b), "--onceki", str(o), "--cikti", c]
            self.assertEqual(main.main(arg + ["--kapasite", "20", "--gecen-yil", str(o), "--butce", str(o)]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(["--bugun", str(b), "--cikti", c]), 2)
            self.assertEqual(main.main(arg + ["--kapasite", "0"]), 2)
            self.assertEqual(main.main(arg + ["--butce", str(Path(tmp) / "yok.csv")]), 1)
            o.write_text("Tarih;Gelir\n01.11.2026;5\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            o.write_text("Tarih;Oda\nx;5\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
