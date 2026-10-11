"""
Zaman Etüdü ve Standart Süre — Workers / Workless kod bloğu
Endüstri Mühendisliği ve Süreç Geliştirme › Endüstri Mühendisi

Kronometre ölçümlerinden performans ve tolerans payıyla operasyon standart sürelerini hesaplar.
  Normal süre (gözlem)   = gözlenen süre × performans % ÷ 100
  Eleman normal süresi   = gözlemlerin normal sürelerinin ortalaması × frekans      (frekans: 1 = her parçada; 1/10 = on parçada bir)
  Standart süre          = normal süre × (1 + tolerans % ÷ 100)
  Saatlik standart çıktı = 3600 ÷ operasyon standart süresi (sn)
  Gerekli gözlem sayısı  = (k ÷ e)² × (n·Σx² − (Σx)²) ÷ (Σx)²        (varsayılan k = 2, e = 0,05: yaklaşık %95 güven, ±%5 hata → çarpan 1600)
Süreler saniyedir; virgül ve nokta ondalık ayırıcıdır. Aykırı gözlemler otomatik atılmaz. Performans değerlendirmesi etüdü yapanın yargısıdır.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"operasyon": ("operasyon", "operasyon adi", "is"), "eleman": ("eleman", "is elemani", "adim"), "sure": ("sure sn", "sure", "gozlenen sure", "olcum sn"),
         "performans": ("performans", "performans yuzde", "tempo", "derecelendirme"), "frekans": ("frekans", "siklik"), "tolerans": ("tolerans", "tolerans yuzde", "tolerans payi", "pay")}
Z, BIR = Decimal(0), Decimal(1)


def ondalik(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-/]", "", str(x)).replace(",", ".")
    try:
        if "/" in s:
            a, b = s.split("/")
            return Decimal(a) / Decimal(b)
        return Decimal(s)
    except (InvalidOperation, ValueError, ZeroDivisionError):
        return None


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("eleman", "sure")):
        if not metin(r["eleman"]):
            continue
        kim = f"{metin(r['operasyon'])} · {metin(r['eleman'])} (satır {r['_satir']})".strip(" ·")
        s, p, f, t = (ondalik(r[a]) for a in ("sure", "performans", "frekans", "tolerans"))
        if s is None or s <= 0:
            uy.append(uyari("Yüksek", "Süre okunamadı", kim, f"'{metin(r['sure'])}'. Gözlem alınmadı."))
            continue
        for v, ad, alt, ust in ((p, "performans", 1, 200), (f, "frekans", 0, 1000), (t, "tolerans", -1, 100)):
            if metin(r[ad]) and (v is None or not alt < v <= ust if ad != "tolerans" else v is None or not 0 <= v <= ust):
                uy.append(uyari("Yüksek", f"{ad.capitalize()} okunamadı", kim, f"'{metin(r[ad])}'. Varsayılan kullanıldı."))
                if ad == "performans":
                    p = None
                elif ad == "frekans":
                    f = None
                else:
                    t = None
        l.append({"operasyon": metin(r["operasyon"]) or "Operasyon", "eleman": metin(r["eleman"]), "sure": s, "performans": p, "frekans": f, "tolerans": t})
    return l, uy


def hesapla(gozlemler: list[dict], tolerans: Decimal = Z, performans: Decimal = Decimal(100), k: Decimal = Decimal(2), hata: Decimal = Decimal(5)) -> dict:
    if not gozlemler:
        raise ValueError("Hesaplanacak gözlem yok.")
    uy, ops = [], {}
    for g in gozlemler:
        ops.setdefault(g["operasyon"], {}).setdefault(g["eleman"], []).append(g)
    carpan = (k / (hata / 100)) ** 2
    operasyonlar = []
    for op, elemanlar in ops.items():
        el = []
        for ad, l in elemanlar.items():
            x = [g["sure"] for g in l]
            n, sx, sxx = len(x), sum(x), sum(v * v for v in x)
            frekans = next((g["frekans"] for g in l if g["frekans"] is not None), BIR)
            tol = next((g["tolerans"] for g in l if g["tolerans"] is not None), tolerans)
            perf = [g["performans"] if g["performans"] is not None else performans for g in l]
            normal = sum((v * p / 100 for v, p in zip(x, perf)), Z) / n
            gerekli = int((carpan * (n * sxx - sx * sx) / (sx * sx)).to_integral_value(rounding="ROUND_CEILING")) if n > 1 else None
            e = {"eleman": ad, "n": n, "ort": sx / n, "min": min(x), "max": max(x), "performans": sum(perf, Z) / n, "frekans": frekans, "tolerans": tol, "normal": normal * frekans,
                 "standart": normal * frekans * (1 + tol / 100), "gerekli": gerekli, "yeterli": None if gerekli is None else n >= gerekli}
            if n == 1:
                uy.append(uyari("Orta", "Tek gözlem", f"{op} · {ad}", "Eleman tek gözlemle hesaplandı; gerekli gözlem sayısı hesaplanamaz."))
            elif not e["yeterli"]:
                uy.append(uyari("Orta", "Gözlem sayısı yetersiz", f"{op} · {ad}", f"{n} gözlem var; ±%{hata.normalize():f} hata için yaklaşık {gerekli} gözlem gerekir (ölçümler çok değişken)."))
            el.append(e)
        ns, ss = sum((e["normal"] for e in el), Z), sum((e["standart"] for e in el), Z)
        operasyonlar.append({"operasyon": op, "elemanlar": el, "gozlenen": sum((e["ort"] * e["frekans"] for e in el), Z), "normal": ns, "standart": ss, "saatlik": Decimal(3600) / ss,
                             "dakika": ss / 60, "yetersiz": sum(1 for e in el if e["yeterli"] is False)})
    return {"operasyonlar": operasyonlar, "uyarilar": uy, "carpan": carpan}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict], vardiya_dk: Decimal | None = None) -> None:
    wb = Workbook()
    y = lambda v: yuvarla(v) if v is not None else None   # noqa: E731
    ws = wb.active
    ws.title = "Standart Süreler"
    _baslik(ws, ["Operasyon", "Eleman Sayısı", "Gözlenen Süre (sn)", "Normal Süre (sn)", "Standart Süre (sn)", "Standart Süre (dk)", "Saatlik Standart Çıktı (adet)",
                 "Vardiya Çıktısı (adet)" if vardiya_dk else "Vardiya Çıktısı (--vardiya-dk verin)", "Gözlemi Yetersiz Eleman"], (30, 9, 13, 13, 13, 13, 14, 16, 12))
    for o in r["operasyonlar"]:
        _satir(ws, [o["operasyon"], len(o["elemanlar"]), y(o["gozlenen"]), y(o["normal"]), y(o["standart"]), yuvarla(o["dakika"], "0.0001"), y(o["saatlik"]),
                    int(vardiya_dk * 60 / o["standart"]) if vardiya_dk else None, o["yetersiz"] or None], para_s=(3, 4, 5, 7), dolgu={9: SARI if o["yetersiz"] else None})
        ws.cell(ws.max_row, 6).number_format = "0.0000"
    ws.append([])
    ws.append(["Standart süre = normal süre × (1 + tolerans). Vardiya çıktısı aşağı yuvarlanır. Aykırı gözlemler atılmamıştır."])
    ws = wb.create_sheet("Elemanlar")
    _baslik(ws, ["Operasyon", "Eleman", "Gözlem Sayısı", "Ortalama (sn)", "En Kısa", "En Uzun", "Performans %", "Frekans", "Tolerans %", "Normal Süre (sn)", "Standart Süre (sn)", "Gerekli Gözlem",
                 "Gözlem Yeterli mi"], (26, 34, 9, 12, 9, 9, 11, 9, 10, 13, 13, 10, 12))
    for o in r["operasyonlar"]:
        for e in o["elemanlar"]:
            _satir(ws, [o["operasyon"], e["eleman"], e["n"], y(e["ort"]), e["min"], e["max"], y(e["performans"]), yuvarla(e["frekans"], "0.0001"), e["tolerans"], y(e["normal"]), y(e["standart"]),
                        e["gerekli"], "" if e["yeterli"] is None else "Evet" if e["yeterli"] else "Hayır"], para_s=(4, 5, 6, 7, 10, 11), dolgu={13: SARI if e["yeterli"] is False else None})
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, vardiya_dk: Decimal | None = None, **secenek) -> dict:
    gozlemler, uy = oku(yol)
    r = hesapla(gozlemler, **secenek)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, r["uyarilar"], vardiya_dk)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Kronometre ölçümlerinden normal ve standart süreyi hesaplar.")
    p.add_argument("--gozlemler", type=Path, default=ORNEK / "gozlemler.csv", help="Operasyon, Eleman, Süre (sn), Performans %%, Frekans, Tolerans %%")
    p.add_argument("--tolerans", help="Dosyada tolerans yoksa kullanılacak pay %% (varsayılan 0; örnek veride 12)")
    p.add_argument("--performans", default="100", help="Dosyada performans yoksa kullanılacak değer %% (varsayılan 100)")
    p.add_argument("--hata", default="5", help="Gerekli gözlem hesabında kabul edilen hata %% (varsayılan 5)")
    p.add_argument("--guven-k", default="2", help="Gerekli gözlem hesabında güven katsayısı (varsayılan 2 ≈ %%95)")
    p.add_argument("--vardiya-dk", help="Net vardiya süresi (dk); verilirse vardiya çıktısı hesaplanır")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "standart_sure.xlsx")
    a = p.parse_args(argv)
    if not a.gozlemler.exists():
        print(f"[X] Dosya bulunamadı: {a.gozlemler}")
        return 1
    tol = para(a.tolerans) if a.tolerans else (Decimal(12) if a.gozlemler == ORNEK / "gozlemler.csv" else Z)
    perf, hata, k, vd = para(a.performans), para(a.hata), para(a.guven_k), para(a.vardiya_dk) if a.vardiya_dk else None
    if tol is None or not 0 <= tol <= 100 or perf is None or not 0 < perf <= 200 or hata is None or not 0 < hata <= 50 or k is None or k <= 0 or (a.vardiya_dk and (vd is None or vd <= 0)):
        print("[X] --tolerans 0-100, --performans 1-200, --hata 0-50 arasında; --guven-k ve --vardiya-dk pozitif olmalı")
        return 2
    try:
        r = calistir(a.gozlemler, a.cikti, vd, tolerans=tol, performans=perf, k=k, hata=hata)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    for o in r["operasyonlar"]:
        print(f"[OK] {o['operasyon']}: normal {yuvarla(o['normal'])} sn · standart {yuvarla(o['standart'])} sn · saatlik {yuvarla(o['saatlik'])} adet"
              + (f" · {o['yetersiz']} elemanda gözlem yetersiz" if o["yetersiz"] else ""))
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
