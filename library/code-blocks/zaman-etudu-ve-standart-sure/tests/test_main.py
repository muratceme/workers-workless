import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
Y = main.yuvarla


def g(eleman, sureler, performans=None, frekans=None, tolerans=None, op="Op"):
    return [{"operasyon": op, "eleman": eleman, "sure": D(str(s)), "performans": None if performans is None else D(performans), "frekans": None if frekans is None else D(str(frekans)),
             "tolerans": None if tolerans is None else D(tolerans)} for s in sureler]


E1, E2, E3, E4 = g("Al", [5.0, 5.2, 4.8, 5.1, 4.9], 100), g("Vidala", [12.0, 12.5, 11.5, 12.2, 11.8], 110), g("Kontrol", [4.0, 4.2, 3.8, 4.1, 3.9], 90), g("Kutu", [20, 22, 18], 100, "0.1")


class HesapTesti(unittest.TestCase):
    def test_eleman_ve_operasyon(self):
        o = main.hesapla(E1 + E2 + E3 + E4, tolerans=D(12))["operasyonlar"][0]
        e = {x["eleman"]: x for x in o["elemanlar"]}
        self.assertEqual([(Y(e[k]["ort"]), Y(e[k]["normal"]), Y(e[k]["standart"])) for k in ("Al", "Vidala", "Kontrol", "Kutu")],
                         [(D("5.00"), D("5.00"), D("5.60")), (D("12.00"), D("13.20"), D("14.78")), (D("4.00"), D("3.60"), D("4.03")), (D("20.00"), D("2.00"), D("2.24"))])
        self.assertEqual((Y(o["gozlenen"]), Y(o["normal"]), Y(o["standart"]), Y(o["saatlik"]), Y(o["dakika"], "0.0001")), (D("23.00"), D("23.80"), D("26.66"), D("135.05"), D("0.4443")))

    def test_gerekli_gozlem(self):
        r = main.hesapla(E1 + E2 + E3 + E4)
        e = {x["eleman"]: (x["gerekli"], x["yeterli"]) for x in r["operasyonlar"][0]["elemanlar"]}
        self.assertEqual(e, {"Al": (2, True), "Vidala": (2, True), "Kontrol": (2, True), "Kutu": (11, False)})       # Kutu: 1600 × 24 ÷ 3600 = 10,67
        self.assertEqual(([u["tur"] for u in r["uyarilar"]], r["operasyonlar"][0]["yetersiz"]), (["Gözlem sayısı yetersiz"], 1))
        self.assertEqual(main.hesapla(E4, hata=D(10))["operasyonlar"][0]["elemanlar"][0]["gerekli"], 3)             # 400 × 24 ÷ 3600 = 2,67
        r = main.hesapla(g("Tek", [7]))
        self.assertEqual((r["operasyonlar"][0]["elemanlar"][0]["gerekli"], [u["tur"] for u in r["uyarilar"]]), (None, ["Tek gözlem"]))
        self.assertEqual(main.hesapla(g("Sabit", [3, 3, 3]))["operasyonlar"][0]["elemanlar"][0]["gerekli"], 0)

    def test_varsayilanlar_ve_gozlem_bazli_performans(self):
        o = main.hesapla(g("A", [10, 10]), tolerans=D(15), performans=D(120))["operasyonlar"][0]
        self.assertEqual((Y(o["normal"]), Y(o["standart"])), (D("12.00"), D("13.80")))
        l = g("A", [10], 80) + g("A", [20], 120)                                                                    # (10 × 0,8 + 20 × 1,2) ÷ 2
        self.assertEqual(Y(main.hesapla(l)["operasyonlar"][0]["normal"]), D("16.00"))
        o = main.hesapla(g("A", [10, 10], tolerans=20) + g("B", [5, 5]), tolerans=D(10))["operasyonlar"][0]
        self.assertEqual(Y(o["standart"]), D("17.50"))                                                              # 10 × 1,20 + 5 × 1,10
        self.assertEqual([x["operasyon"] for x in main.hesapla(g("A", [1, 1]) + g("A", [2, 2], op="Op2"))["operasyonlar"]], ["Op", "Op2"])
        with self.assertRaises(ValueError):
            main.hesapla([])

    def test_sayi_okuma(self):
        self.assertEqual([main.ondalik(x) for x in ("12,5", "12.5", "1.250", "1/10", "1/0", "%110", "", "x")], [D("12.5"), D("12.5"), D("1.250"), D("0.1"), None, D(110), None, None])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "s.xlsx"
            r = main.calistir(ORNEK / "gozlemler.csv", cikti, D(450), tolerans=D(12))
            o = r["operasyonlar"][0]
            self.assertEqual((o["operasyon"], len(o["elemanlar"]), Y(o["normal"]), Y(o["standart"]), Y(o["saatlik"]), o["yetersiz"]), ("Montaj A", 4, D("23.80"), D("26.66"), D("135.05"), 1))
            self.assertEqual([u["tur"] for u in r["uyarilar"]], ["Gözlem sayısı yetersiz"])
            wb = load_workbook(cikti)
            self.assertEqual((wb.sheetnames, wb["Standart Süreler"]["H2"].value), (["Standart Süreler", "Elemanlar", "Uyarılar"], 1012))      # 450 × 60 ÷ 26,656

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "g.csv"
            y.write_text("Eleman;Süre (sn);Performans %;Frekans;Tolerans %\nA;10,5;250;0;150\nA;0;;;\nB;yok;;;\nA;9.5;100;1/2;10\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual([(x["sure"], x["performans"], x["frekans"], x["tolerans"], x["operasyon"]) for x in l], [(D("10.5"), None, None, None, "Operasyon"), (D("9.5"), D(100), D("0.5"), D(10), "Operasyon")])
            self.assertEqual(sorted(u["tur"] for u in uy), ["Frekans okunamadı", "Performans okunamadı", "Süre okunamadı", "Süre okunamadı", "Tolerans okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--gozlemler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tolerans", "10", "--vardiya-dk", "450", "--hata", "10"]), 0)
            self.assertEqual(main.main(arg + ["--tolerans", "150"]), 2)
            self.assertEqual(main.main(arg + ["--vardiya-dk", "0"]), 2)
            self.assertEqual(main.main(["--gozlemler", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Eleman;Süre (sn)\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
