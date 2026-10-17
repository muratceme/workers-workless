"""
Kumaş Fiyat Takibi — Gerisi Bende kod bloğu
Kumaş ve Aksesuar Satın Alma › Kumaş Satın Alma Uzmanı

Tedarikçi kumaş ve iplik fiyat listelerini bir araya getirir; kalite bazında fiyat geçmişi ve güncel teklif karşılaştırması üretir.
  Fiyat serisi         = kalite + tedarikçi + birim + para birimi;   aynı serinin kayıtları tarihe göre sıralanır
  Son değişim %        = (son fiyat − önceki fiyat) ÷ önceki fiyat × 100
  Dönem değişimi %     = (son fiyat − ilk fiyat) ÷ ilk fiyat × 100
  Güncel teklif        = her serinin en son tarihli fiyatı
  Kalite karşılaştırması = aynı kalite, birim ve para birimindeki güncel teklifler: en düşük, ortanca, en yüksek;   fark % = (en yüksek − en düşük) ÷ en düşük × 100
  Uyarı                = son değişim --artis (%10) üstünde artış · en son fiyat --eski (90) günden eski
Farklı para birimleri ve birimler birbirine çevrilmez; ayrı karşılaştırılır. Fiyatın kapsamını (en, gramaj, teslim şekli, vade) kod bilmez.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"tarih": ("tarih", "fiyat tarihi", "teklif tarihi", "gecerlilik tarihi"), "tedarikci": ("tedarikci", "tedarikci adi", "firma"), "kalite": ("kalite", "kumas kalitesi", "kumas", "iplik", "urun", "kalite kodu"),
         "birim": ("birim", "olcu birimi"), "para": ("para birimi", "doviz", "doviz cinsi"), "fiyat": ("fiyat", "birim fiyat", "teklif fiyati")}


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("tarih", "tedarikci", "kalite", "fiyat")):
        ted, kalite = metin(r["tedarikci"]), metin(r["kalite"])
        if not ted and not kalite:
            continue
        t, fiyat = tarih(r["tarih"]), para(r["fiyat"])
        if t is None or fiyat is None or fiyat <= 0 or not ted or not kalite:
            uy.append(uyari("Yüksek", "Fiyat satırı okunamadı", f"{kalite} · {ted}", f"Satır {r['_satir']}: tarih '{metin(r['tarih'])}', fiyat '{metin(r['fiyat'])}', tedarikçi ya da kalite okunamadı. Satır alınmadı."))
            continue
        l.append({"tarih": t, "tedarikci": ted, "kalite": kalite, "birim": metin(r["birim"]) or "m", "para": (metin(r["para"]) or "TRY").upper(), "fiyat": fiyat, "satir": r["_satir"]})
    return l, uy


def _oran(pay: Decimal, payda: Decimal) -> Decimal | None:
    return yuvarla(pay * 100 / payda) if payda else None


def _medyan(d: list[Decimal]) -> Decimal:
    s, n = sorted(d), len(d)
    return s[n // 2] if n % 2 else (s[n // 2 - 1] + s[n // 2]) / 2


def takip(fiyatlar: list[dict], bugun: date, artis: Decimal = Decimal(10), eski: int = 90) -> dict:
    if not fiyatlar:
        raise ValueError("İzlenecek fiyat kaydı yok.")
    uy, grup = [], {}
    for f in fiyatlar:
        grup.setdefault((katla(f["kalite"]), katla(f["birim"]), f["para"], katla(f["tedarikci"])), []).append(f)
    seriler, gecmis = [], []
    for g in grup.values():
        g.sort(key=lambda f: (f["tarih"], f["satir"]))
        kim = f"{g[-1]['kalite']} · {g[-1]['tedarikci']}"
        temiz = []
        for f in g:
            if temiz and temiz[-1]["tarih"] == f["tarih"]:
                uy.append(uyari("Bilgi", "Aynı tarihte iki fiyat", kim, f"{f['tarih']:%d.%m.%Y} için birden çok fiyat var; dosyada sonra gelen ({f['fiyat']}) kullanıldı."))
                temiz[-1] = f
            else:
                temiz.append(f)
        for i, f in enumerate(temiz):
            f["degisim"] = _oran(f["fiyat"] - temiz[i - 1]["fiyat"], temiz[i - 1]["fiyat"]) if i else None
            gecmis.append(f)
        ilk, son, fy = temiz[0], temiz[-1], [f["fiyat"] for f in temiz]
        s = {"kalite": son["kalite"], "tedarikci": son["tedarikci"], "birim": son["birim"], "para": son["para"], "kayit": len(temiz), "ilk_tarih": ilk["tarih"], "ilk": ilk["fiyat"], "onceki": temiz[-2]["fiyat"] if len(temiz) > 1 else None,
             "son_tarih": son["tarih"], "son": son["fiyat"], "son_degisim": son["degisim"], "donem": _oran(son["fiyat"] - ilk["fiyat"], ilk["fiyat"]) if len(temiz) > 1 else None, "enk": min(fy), "eny": max(fy),
             "yas": (bugun - son["tarih"]).days}
        seriler.append(s)
        if s["son_degisim"] is not None and s["son_degisim"] > artis:
            uy.append(uyari("Orta", "Fiyat artışı", kim, f"{s['onceki']} → {s['son']} {s['para']}/{s['birim']} (%{s['son_degisim']}; {s['son_tarih']:%d.%m.%Y}). Dönem başından değişim %{s['donem']}."))
        elif s["son_degisim"] is not None and s["son_degisim"] < -artis:
            uy.append(uyari("Bilgi", "Fiyat düşüşü", kim, f"{s['onceki']} → {s['son']} {s['para']}/{s['birim']} (%{s['son_degisim']}; {s['son_tarih']:%d.%m.%Y})."))
        if s["yas"] > eski:
            uy.append(uyari("Bilgi", "Fiyat eski", kim, f"Son fiyat {s['son_tarih']:%d.%m.%Y} tarihli ({s['yas']} gün önce). Güncel teklif isteyin."))
    seriler.sort(key=lambda s: (katla(s["kalite"]), katla(s["birim"]), s["para"], s["son"], katla(s["tedarikci"])))
    kg = {}
    for s in seriler:
        kg.setdefault((katla(s["kalite"]), katla(s["birim"]), s["para"]), []).append(s)
    kaliteler = []
    for g in kg.values():
        enk, eny = g[0], g[-1]
        for s in g:
            s["sira"], s["enk_fark"] = g.index(s) + 1, _oran(s["son"] - enk["son"], enk["son"])
        kaliteler.append({"kalite": enk["kalite"], "birim": enk["birim"], "para": enk["para"], "teklif": len(g), "enk": enk["son"], "enk_ted": enk["tedarikci"], "ortanca": yuvarla(_medyan([s["son"] for s in g]), "0.0001").normalize(),
                          "eny": eny["son"], "eny_ted": eny["tedarikci"], "fark": _oran(eny["son"] - enk["son"], enk["son"]) if len(g) > 1 else None, "en_eski": max(s["yas"] for s in g)})
    return {"seriler": seriler, "kaliteler": kaliteler, "gecmis": sorted(gecmis, key=lambda f: (katla(f["kalite"]), katla(f["birim"]), f["para"], katla(f["tedarikci"]), f["tarih"])),
            "ozet": {"kayit": len(fiyatlar), "seri": len(seriler), "kalite": len({katla(s["kalite"]) for s in seriler}), "tedarikci": len({katla(s["tedarikci"]) for s in seriler}),
                     "artan": sum(1 for s in seriler if s["son_degisim"] is not None and s["son_degisim"] > artis), "eski": sum(1 for s in seriler if s["yas"] > eski)}, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, bugun: date, artis: Decimal, eski: int) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kalite Karşılaştırması"
    _baslik(ws, ["Kalite", "Birim", "Para Birimi", "Güncel Teklif", "En Düşük Fiyat", "En Düşük Fiyatlı Tedarikçi", "Ortanca", "En Yüksek Fiyat", "En Yüksek Fiyatlı Tedarikçi", "Fark %", "En Eski Teklifin Yaşı (gün)"],
            (30, 8, 9, 9, 13, 28, 12, 13, 28, 9, 12))
    for k in r["kaliteler"]:
        _satir(ws, [k["kalite"], k["birim"], k["para"], k["teklif"], k["enk"], k["enk_ted"], k["ortanca"], k["eny"], k["eny_ted"], k["fark"], k["en_eski"]], para_s=(5, 7, 8, 10), dolgu={11: SARI if k["en_eski"] > eski else None})
    ws.append([])
    ws.append([f"Durum tarihi: {bugun:%d.%m.%Y}. Güncel teklif = her tedarikçinin en son tarihli fiyatı; fiyatlar para birimine çevrilmez."])
    ws = wb.create_sheet("Güncel Fiyatlar")
    _baslik(ws, ["Kalite", "Birim", "Para Birimi", "Sıra", "Tedarikçi", "Son Fiyat", "En Düşüğe Göre Fark %", "Son Fiyat Tarihi", "Yaş (gün)", "Önceki Fiyat", "Son Değişim %", "İlk Fiyat", "İlk Fiyat Tarihi", "Dönem Değişimi %", "En Düşük", "En Yüksek",
                 "Kayıt"], (30, 8, 9, 6, 28, 12, 11, 12, 9, 12, 10, 12, 12, 11, 12, 12, 7))
    for s in r["seriler"]:
        _satir(ws, [s["kalite"], s["birim"], s["para"], s["sira"], s["tedarikci"], s["son"], s["enk_fark"], s["son_tarih"], s["yas"], s["onceki"], s["son_degisim"], s["ilk"], s["ilk_tarih"], s["donem"], s["enk"], s["eny"], s["kayit"]],
               para_s=(6, 7, 10, 11, 12, 14, 15, 16), tarih_s=(8, 13),
               dolgu={4: YESIL if s["sira"] == 1 else None, 9: SARI if s["yas"] > eski else None, 11: None if s["son_degisim"] is None else KIRMIZI if s["son_degisim"] > artis else YESIL if s["son_degisim"] < 0 else None})
    ws = wb.create_sheet("Fiyat Geçmişi")
    _baslik(ws, ["Kalite", "Tedarikçi", "Birim", "Para Birimi", "Tarih", "Fiyat", "Önceki Kayda Göre %"], (30, 28, 8, 9, 12, 12, 12))
    for f in r["gecmis"]:
        _satir(ws, [f["kalite"], f["tedarikci"], f["birim"], f["para"], f["tarih"], f["fiyat"], f["degisim"]], para_s=(6, 7), tarih_s=(5,))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, artis: Decimal = Decimal(10), eski: int = 90) -> dict:
    fiyatlar, uy = oku(yol)
    r = takip(fiyatlar, bugun, artis, eski)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, bugun, artis, eski)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Tedarikçi kumaş ve iplik fiyat listelerini toplayıp kalite bazında fiyat geçmişi tablosu oluşturur.")
    p.add_argument("--fiyatlar", type=Path, help="Tarih, Tedarikçi, Kalite, Birim, Para Birimi, Fiyat")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--artis", default="10", help="Son fiyat değişimi bu yüzdeyi aşarsa uyarılır (varsayılan 10)")
    p.add_argument("--eski", type=int, default=90, help="Son fiyatı bu kadar günden eski olan teklifler bildirilir (varsayılan 90)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kumas_fiyat_takibi.xlsx")
    a = p.parse_args(argv)
    ornek = a.fiyatlar is None
    yol = ORNEK / "fiyat_listeleri.csv" if ornek else a.fiyatlar
    bugun, artis = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today(), para(a.artis)
    if bugun is None or artis is None or artis < 0 or a.eski < 1:
        print("[X] --tarih GG.AA.YYYY; --artis eksi olamaz; --eski en az 1 olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, bugun, artis, a.eski)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {o['kayit']} fiyat kaydı · {o['kalite']} kalite · {o['tedarikci']} tedarikçi · {o['seri']} fiyat serisi · artış uyarısı {o['artan']} · eski fiyat {o['eski']}")
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
