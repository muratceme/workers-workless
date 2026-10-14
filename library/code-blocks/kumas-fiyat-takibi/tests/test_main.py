import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)
_n = [0]


def f(kalite, ted, t, fiyat, birim="m", para="TRY"):
    _n[0] += 1
    return {"tarih": t, "tedarikci": ted, "kalite": kalite, "birim": birim, "para": para, "fiyat": D(str(fiyat)), "satir": _n[0]}


def veri():
    return [f("K1", "T1", date(2026, 8, 1), 110), f("K1", "T1", date(2026, 6, 1), 100), f("k1", "t1", date(2026, 10, 1), 125), f("K1", "T2", date(2026, 9, 15), 118), f("K1", "T3", date(2026, 7, 1), 3, para="USD"),
            f("K2", "T1", date(2026, 5, 1), 200, "kg"), f("K2", "T1", date(2026, 9, 1), 180, "kg")]


class TakipTesti(unittest.TestCase):
    def test_seriler(self):
        r = main.takip(veri(), BUGUN)
        self.assertEqual([(s["kalite"], s["tedarikci"], s["para"], s["sira"], s["kayit"], s["ilk"], s["onceki"], s["son"], s["son_degisim"], s["donem"], s["enk"], s["eny"], s["yas"], s["enk_fark"]) for s in r["seriler"]],
                         [("K1", "T2", "TRY", 1, 1, D(118), None, D(118), None, None, D(118), D(118), 27, D("0.00")), ("k1", "t1", "TRY", 2, 3, D(100), D(110), D(125), D("13.64"), D("25.00"), D(100), D(125), 11, D("5.93")),
                          ("K1", "T3", "USD", 1, 1, D(3), None, D(3), None, None, D(3), D(3), 103, D("0.00")), ("K2", "T1", "TRY", 1, 2, D(200), D(200), D(180), D("-10.00"), D("-10.00"), D(180), D(200), 41, D("0.00"))])
        self.assertEqual([(k["kalite"], k["birim"], k["para"], k["teklif"], k["enk"], k["enk_ted"], k["ortanca"], k["eny"], k["eny_ted"], k["fark"], k["en_eski"]) for k in r["kaliteler"]],
                         [("K1", "m", "TRY", 2, D(118), "T2", D("121.5"), D(125), "t1", D("5.93"), 27), ("K1", "m", "USD", 1, D(3), "T3", D(3), D(3), "T3", None, 103), ("K2", "kg", "TRY", 1, D(180), "T1", D(180), D(180), "T1", None, 41)])
        self.assertEqual([(x["tedarikci"], x["tarih"].month, x["fiyat"], x["degisim"]) for x in r["gecmis"] if x["kalite"].upper() == "K1" and x["para"] == "TRY"],
                         [("T1", 6, D(100), None), ("T1", 8, D(110), D("10.00")), ("t1", 10, D(125), D("13.64")), ("T2", 9, D(118), None)])
        self.assertEqual((r["ozet"], sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"])),
                         ({"kayit": 7, "seri": 4, "kalite": 2, "tedarikci": 3, "artan": 1, "eski": 1}, [("Bilgi", "Fiyat eski", "K1 · T3"), ("Orta", "Fiyat artışı", "k1 · t1")]))

    def test_esikler(self):
        u = lambda **kw: sorted(x["tur"] for x in main.takip(veri(), BUGUN, **kw)["uyarilar"])   # noqa: E731
        self.assertEqual((u(artis=D("13.64"), eski=103), u(artis=D("9.99"), eski=40)), ([], ["Fiyat artışı", "Fiyat düşüşü", "Fiyat eski", "Fiyat eski"]))
        r = main.takip([f("K", "T", date(2026, 9, 1), 10), f("K", "T", date(2026, 9, 1), 12), f("K", "T", date(2026, 8, 1), 8)], BUGUN)
        self.assertEqual(([(s["kayit"], s["onceki"], s["son"], s["son_degisim"]) for s in r["seriler"]], [u["tur"] for u in r["uyarilar"]]), ([(2, D(8), D(12), D("50.00"))], ["Aynı tarihte iki fiyat", "Fiyat artışı"]))
        with self.assertRaises(ValueError):
            main.takip([], BUGUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "fiyat_listeleri.csv", cikti, BUGUN)
            self.assertEqual([(k["kalite"], k["para"], k["teklif"], k["enk"], k["enk_ted"], k["ortanca"], k["eny"], k["eny_ted"], k["fark"]) for k in r["kaliteler"]],
                             [("30/1 Süprem %100 pamuk", "TRY", 3, D(310), "Deneme Tekstil Ltd. Şti.", D(320), D(345), "Örnek Örme A.Ş.", D("11.29")),
                              ("İki iplik şardonlu", "TRY", 2, D(365), "Numune Kumaş San.", D("371.5"), D(378), "Örnek Örme A.Ş.", D("3.56")),
                              ("Poplin 40/1", "USD", 2, D("2.10"), "Kurgu Dokuma A.Ş.", D("2.175"), D("2.25"), "Deneme Tekstil Ltd. Şti.", D("7.14"))])
            self.assertEqual([(s["tedarikci"][:6], s["son"], s["son_degisim"], s["donem"], s["yas"]) for s in r["seriler"]],
                             [("Deneme", D(310), D("6.90"), D("6.90"), 53), ("Numune", D(320), None, None, 160), ("Örnek ", D(345), D("15.00"), D("21.05"), 27), ("Numune", D(365), None, None, 48), ("Örnek ", D(378), D("5.00"), D("5.00"), 41),
                              ("Kurgu ", D("2.10"), D("-12.50"), D("-12.50"), 20), ("Deneme", D("2.25"), None, None, 12)])
            self.assertEqual((r["ozet"], [(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]]),
                             ({"kayit": 12, "seri": 7, "kalite": 3, "tedarikci": 4, "artan": 1, "eski": 1},
                              [("Orta", "Fiyat artışı", "30/1 Süprem %100 pamuk · Örnek Örme A.Ş."), ("Bilgi", "Fiyat düşüşü", "Poplin 40/1 · Kurgu Dokuma A.Ş."), ("Bilgi", "Fiyat eski", "30/1 Süprem %100 pamuk · Numune Kumaş San.")]))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Kalite Karşılaştırması", "Güncel Fiyatlar", "Fiyat Geçmişi", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "f.csv"
            y.write_text("Teklif Tarihi;Firma;Kumaş;Döviz;Teklif Fiyatı\n01.09.2026;A;Süprem;usd;2,45\n02.09.2026;A;Süprem;;1.250,50\nx;A;Süprem;;5\n03.09.2026;A;Süprem;;0\n03.09.2026;;Süprem;;5\n03.09.2026;A;;;5\n;;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["tarih"], x["tedarikci"], x["kalite"], x["birim"], x["para"], x["fiyat"]) for x in l], [u["tur"] for u in uy]),
                             ([(date(2026, 9, 1), "A", "Süprem", "m", "USD", D("2.45")), (date(2026, 9, 2), "A", "Süprem", "m", "TRY", D("1250.50"))], ["Fiyat satırı okunamadı"] * 4))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--fiyatlar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "12.10.2026", "--artis", "5", "--eski", "60"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--artis", "-1"], ["--eski", "0"])], [2] * 3)
            self.assertEqual(main.main(["--fiyatlar", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Tarih;Tedarikçi;Kalite\n01.09.2026;A;B\n", "Tarih;Tedarikçi;Kalite;Fiyat\n;;;\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
