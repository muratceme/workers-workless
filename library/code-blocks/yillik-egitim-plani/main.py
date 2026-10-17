"""
Yıllık Eğitim Planı — Gerisi Bende kod bloğu
İnsan Kaynakları › Eğitim ve Gelişim Uzmanı

Eğitim ihtiyaç listesinden bütçeye sığan yıllık eğitim planını, aylık takvimi ve departman özetini oluşturur.
  Eğitim maliyeti   = katılımcı sayısı × kişi başı maliyet + sabit maliyet (eğitmen, salon vb.)
  Adam-saat         = katılımcı sayısı × süre (saat)
  Seçim sırası      = önce zorunlu eğitimler; sonra öncelik (1 en yüksek), eşitlikte düşük maliyetli olan
  Bütçe             = zorunlu eğitimler her durumda plana girer; diğerleri kalan bütçe yettiği sürece (yetmeyen atlanır, sıradaki daha ucuz eğitim denenir)
  Ay ataması        = "Tercih Ay" doluysa o ay; boşsa o ana kadar adam-saat yükü en az olan ay (eşitlikte yıla yayacak sırayla); --kapali-aylar hariç
  Özetler           = ay ve departman bazında eğitim sayısı, katılımcı, adam-saat ve maliyet
Öncelik ve zorunluluk sizin girdinizdir; kod hangi eğitimin yasal olarak zorunlu olduğunu bilmez. Ay ataması bir taslaktır: eğitmen, salon ve üretim takvimine göre siz düzeltirsiniz.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"egitim": ("egitim", "egitim adi", "konu"), "departman": ("departman", "birim", "hedef kitle"), "katilimci": ("katilimci sayisi", "katilimci", "kisi sayisi"), "sure": ("sure", "sure saat", "saat"),
         "kisi": ("kisi basi maliyet", "kisi basi", "birim maliyet"), "sabit": ("sabit maliyet", "sabit", "egitmen ucreti"), "oncelik": ("oncelik",), "zorunlu": ("zorunlu", "yasal zorunlu"), "ay": ("tercih ay", "ay", "planlanan ay")}
AYLAR = ("Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık")
AY_NO = {katla(a): i for i, a in enumerate(AYLAR, 1)}
YAYILIM = (1, 7, 4, 10, 2, 8, 5, 11, 3, 9, 6, 12)
EVET = {"evet", "e", "x", "var", "1", "true", "yes", "zorunlu"}
Z = Decimal(0)


def ay_coz(x) -> int | None:
    s = katla(x)
    if s in AY_NO:
        return AY_NO[s]
    return int(s) if s.isdigit() and 1 <= int(s) <= 12 else None


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("egitim", "katilimci")):
        e = metin(r["egitim"])
        if not e:
            continue
        k, s, kb, sb = para(r["katilimci"]), (para(r["sure"]) if metin(r["sure"]) else Z), (para(r["kisi"]) if metin(r["kisi"]) else Z), (para(r["sabit"]) if metin(r["sabit"]) else Z)
        onc, ay = (para(r["oncelik"]) if metin(r["oncelik"]) else Decimal(3)), (ay_coz(r["ay"]) if metin(r["ay"]) else None)
        if k is None or k <= 0 or k != k.to_integral_value() or None in (s, kb, sb, onc) or s < 0 or kb < 0 or sb < 0 or onc < 1 or (metin(r["ay"]) and ay is None):
            uy.append(uyari("Yüksek", "Eğitim satırı okunamadı", f"{e} (satır {r['_satir']})", f"Katılımcı '{metin(r['katilimci'])}', süre '{metin(r['sure'])}', maliyetler '{metin(r['kisi'])}' / '{metin(r['sabit'])}', öncelik '{metin(r['oncelik'])}', "
                            f"ay '{metin(r['ay'])}'. Katılımcı pozitif tam sayı, ay 1–12 veya ay adı olmalı. Satır plana alınmadı."))
            continue
        l.append({"egitim": e, "departman": metin(r["departman"]) or "Genel", "katilimci": int(k), "sure": s, "kisi": kb, "sabit": sb, "oncelik": int(onc), "zorunlu": katla(r["zorunlu"]) in EVET, "ay": ay})
    return l, uy


def planla(ihtiyaclar: list[dict], butce: Decimal | None = None, kapali: set | None = None) -> dict:
    if not ihtiyaclar:
        raise ValueError("Planlanacak eğitim yok.")
    kapali, uy = kapali or set(), []
    acik = [a for a in YAYILIM if a not in kapali]
    if not acik:
        raise ValueError("Bütün aylar kapalı; eğitim atanacak ay yok.")
    hazir = [e | {"maliyet": e["katilimci"] * e["kisi"] + e["sabit"], "adam_saat": e["katilimci"] * e["sure"]} for e in ihtiyaclar]
    sirali = sorted(hazir, key=lambda e: (not e["zorunlu"], e["oncelik"] if not e["zorunlu"] else 0, e["maliyet"], e["egitim"], e["departman"]))
    harcanan, plan, disi = Z, [], []
    for e in sirali:
        if e["zorunlu"] or butce is None or harcanan + e["maliyet"] <= butce:
            harcanan += e["maliyet"]
            plan.append(e)
        else:
            disi.append(e | {"neden": f"Kalan bütçe {butce - harcanan:,.2f}, eğitim maliyeti {e['maliyet']:,.2f}"})
    zorunlu_top = sum((e["maliyet"] for e in plan if e["zorunlu"]), Z)
    if butce is not None and zorunlu_top > butce:
        uy.append(uyari("Yüksek", "Zorunlu eğitimler bütçeyi aşıyor", "Bütçe", f"Zorunlu eğitimlerin maliyeti {zorunlu_top:,.2f}, bütçe {butce:,.2f}. Zorunlu eğitimler yine de plana alındı; {zorunlu_top - butce:,.2f} ek bütçe gerekir."))
    yuk = {a: Z for a in range(1, 13)}
    for e in plan:
        if e["ay"]:
            yuk[e["ay"]] += e["adam_saat"]
            if e["ay"] in kapali:
                uy.append(uyari("Orta", "Tercih edilen ay kapalı", e["egitim"], f"{AYLAR[e['ay'] - 1]} kapalı aylar arasında ama eğitim için bu ay istenmiş. Tercih korundu; kontrol edin."))
    for e in plan:
        if not e["ay"]:
            e["ay"] = min(acik, key=lambda a: (yuk[a], acik.index(a)))
            e["atanan"] = True
            yuk[e["ay"]] += e["adam_saat"]
    plan.sort(key=lambda e: (e["ay"], not e["zorunlu"], e["oncelik"], e["egitim"], e["departman"]))
    for e in disi:
        uy.append(uyari("Orta" if e["oncelik"] == 1 else "Bilgi", "Bütçe dışı kalan eğitim", f"{e['egitim']} · {e['departman']}", f"Öncelik {e['oncelik']}. {e['neden']}."))

    def ozetle(anahtar) -> list[dict]:
        g = {}
        for e in plan:
            o = g.setdefault(anahtar(e), {"egitim": 0, "katilimci": 0, "adam_saat": Z, "maliyet": Z})
            o["egitim"] += 1
            o["katilimci"] += e["katilimci"]
            o["adam_saat"] += e["adam_saat"]
            o["maliyet"] += e["maliyet"]
        return g

    aylik = ozetle(lambda e: e["ay"])
    dep = ozetle(lambda e: e["departman"])
    toplam = {"egitim": len(plan), "katilimci": sum(e["katilimci"] for e in plan), "adam_saat": sum((e["adam_saat"] for e in plan), Z), "maliyet": harcanan, "butce": butce, "kalan": butce - harcanan if butce is not None else None,
              "disi": len(disi), "disi_maliyet": sum((e["maliyet"] for e in disi), Z), "zorunlu": sum(1 for e in plan if e["zorunlu"]), "zorunlu_maliyet": zorunlu_top}
    return {"plan": plan, "disi": disi, "aylik": [{"ay": a, "ad": AYLAR[a - 1]} | aylik.get(a, {"egitim": 0, "katilimci": 0, "adam_saat": Z, "maliyet": Z}) for a in range(1, 13)],
            "departmanlar": sorted(({"departman": d} | o for d, o in dep.items()), key=lambda o: (-o["maliyet"], o["departman"])), "toplam": toplam, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, yil: int, uy: list[dict]) -> None:
    wb, t = Workbook(), r["toplam"]
    ws = wb.active
    ws.title = "Plan"
    _baslik(ws, ["Ay", "Eğitim", "Departman", "Zorunlu", "Öncelik", "Katılımcı", "Süre (saat)", "Adam-Saat", "Maliyet", "Ay Seçimi", "Tarih / Eğitmen (siz doldurun)"], (10, 36, 20, 9, 8, 10, 10, 10, 15, 12, 30))
    for e in r["plan"]:
        _satir(ws, [AYLAR[e["ay"] - 1], e["egitim"], e["departman"], "Evet" if e["zorunlu"] else "", None if e["zorunlu"] else e["oncelik"], e["katilimci"], e["sure"], e["adam_saat"], e["maliyet"], "Önerildi" if e.get("atanan") else "Tercih", ""],
               para_s=(9,), dolgu={4: SARI if e["zorunlu"] else None, 10: MAVI if e.get("atanan") else None})
    n = _satir(ws, ["Toplam", f"{t['egitim']} eğitim", "", f"{t['zorunlu']} zorunlu", None, t["katilimci"], None, t["adam_saat"], t["maliyet"]], para_s=(9,))
    for c in ws[n]:
        c.font = Font(bold=True)
    ws.append([])
    ws.append([f"{yil} planı. " + (f"Bütçe {t['butce']:,.2f}; planlanan {t['maliyet']:,.2f}; kalan {t['kalan']:,.2f}. " if t["butce"] is not None else "Bütçe verilmedi; bütün eğitimler plana alındı. ")
               + "'Önerildi' yazan aylar adam-saat yükünü dengelemek için kod tarafından atanmıştır; taslaktır."])
    ws = wb.create_sheet("Aylık Özet")
    _baslik(ws, ["Ay", "Eğitim Sayısı", "Katılımcı", "Adam-Saat", "Maliyet"], (12, 10, 10, 12, 16))
    for a in r["aylik"]:
        _satir(ws, [a["ad"], a["egitim"] or None, a["katilimci"] or None, a["adam_saat"] or None, a["maliyet"] or None], para_s=(5,))
    ws = wb.create_sheet("Departman Özeti")
    _baslik(ws, ["Departman", "Eğitim Sayısı", "Katılımcı", "Adam-Saat", "Maliyet", "Bütçedeki Payı %"], (24, 10, 10, 12, 16, 11))
    for d in r["departmanlar"]:
        _satir(ws, [d["departman"], d["egitim"], d["katilimci"], d["adam_saat"], d["maliyet"], yuvarla(d["maliyet"] / t["maliyet"] * 100) if t["maliyet"] else None], para_s=(5, 6))
    ws = wb.create_sheet("Bütçe Dışı")
    _baslik(ws, ["Eğitim", "Departman", "Öncelik", "Katılımcı", "Maliyet", "Neden"], (36, 20, 8, 10, 15, 60))
    for e in r["disi"]:
        _satir(ws, [e["egitim"], e["departman"], e["oncelik"], e["katilimci"], e["maliyet"], e["neden"]], para_s=(5,))
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, yil: int, butce: Decimal | None = None, kapali: set | None = None) -> dict:
    ihtiyaclar, uy = oku(yol)
    r = planla(ihtiyaclar, butce, kapali)
    uy += r["uyarilar"]
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, yil, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Eğitim ihtiyaç listesinden bütçeye sığan yıllık eğitim planı oluşturur.")
    p.add_argument("--ihtiyaclar", type=Path, default=ORNEK / "egitim_ihtiyaclari.csv", help="Eğitim, Departman, Katılımcı Sayısı, Süre (saat), Kişi Başı Maliyet, Sabit Maliyet, Öncelik, Zorunlu, Tercih Ay")
    p.add_argument("--butce", help="Yıllık eğitim bütçesi; verilmezse bütün eğitimler plana alınır")
    p.add_argument("--yil", type=int, help="Plan yılı (varsayılan: gelecek yıl)")
    p.add_argument("--kapali-aylar", default="", help="Eğitim atanmayacak aylar, virgülle (ör. 7,8)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "yillik_egitim_plani.xlsx")
    a = p.parse_args(argv)
    ornek = a.ihtiyaclar == ORNEK / "egitim_ihtiyaclari.csv"
    butce = para(a.butce) if a.butce else Decimal(100000) if ornek else None
    kapali = [ay_coz(x) for x in a.kapali_aylar.split(",") if x.strip()]
    if (a.butce and (butce is None or butce < 0)) or None in kapali or (a.yil is not None and not 2000 <= a.yil <= 2100):
        print("[X] --butce sıfır veya pozitif sayı, --kapali-aylar 1–12 arası ay numaraları veya ay adları, --yil geçerli bir yıl olmalı")
        return 2
    if not a.ihtiyaclar.exists():
        print(f"[X] Dosya bulunamadı: {a.ihtiyaclar}")
        return 1
    try:
        r = calistir(a.ihtiyaclar, a.cikti, a.yil or (2027 if ornek else date.today().year + 1), butce, set(kapali))
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {t['egitim']} eğitim planlandı ({t['zorunlu']} zorunlu) · {t['katilimci']} katılımcı · {t['adam_saat']:,.0f} adam-saat · maliyet {t['maliyet']:,.2f}" + (f" · kalan bütçe {t['kalan']:,.2f}" if t["kalan"] is not None else ""))
    if t["disi"]:
        print(f"[!] {t['disi']} eğitim bütçe dışı kaldı (toplam {t['disi_maliyet']:,.2f})")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
