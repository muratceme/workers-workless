import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def e(ad, dep, kisi, sure, kb=0, sabit=0, oncelik=3, zorunlu=False, ay=None):
    return {"egitim": ad, "departman": dep, "katilimci": kisi, "sure": D(sure), "kisi": D(kb), "sabit": D(sabit), "oncelik": oncelik, "zorunlu": zorunlu, "ay": ay}


IHT = lambda: [e("İSG", "Üretim", 40, 8, 500, zorunlu=True, ay=2), e("Yangın", "Tüm", 60, 4, 0, 15000, zorunlu=True), e("Excel", "Finans", 10, 16, 2000, oncelik=2, ay=3), e("Liderlik", "Yönetim", 8, 16, 5000, 10000, 1),   # noqa: E731
               e("Sunum", "Satış", 12, 8, 1500, oncelik=3), e("İngilizce", "Satış", 5, 40, 3000, oncelik=2)]
ozet = lambda r: [(x["egitim"], x["ay"], x["maliyet"], x["adam_saat"], bool(x.get("atanan"))) for x in r["plan"]]   # noqa: E731


class PlanTesti(unittest.TestCase):
    def test_butceli_plan(self):
        r = main.planla(IHT(), D(100000))
        self.assertEqual(ozet(r), [("Yangın", 1, D(15000), D(240), True), ("İSG", 2, D(20000), D(320), False), ("İngilizce", 4, D(15000), D(200), True), ("Liderlik", 7, D(50000), D(128), True)])
        t = r["toplam"]
        self.assertEqual((t["egitim"], t["katilimci"], t["adam_saat"], t["maliyet"], t["kalan"], t["disi"], t["disi_maliyet"], t["zorunlu"], t["zorunlu_maliyet"]), (4, 113, D(888), D(100000), D(0), 2, D(38000), 2, D(35000)))
        self.assertEqual(([(x["egitim"], x["oncelik"]) for x in r["disi"]], sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"])),
                         ([("Excel", 2), ("Sunum", 3)], [("Bilgi", "Bütçe dışı kalan eğitim", "Excel · Finans"), ("Bilgi", "Bütçe dışı kalan eğitim", "Sunum · Satış")]))
        self.assertEqual([(d["departman"], d["egitim"], d["katilimci"], d["maliyet"]) for d in r["departmanlar"]], [("Yönetim", 1, 8, D(50000)), ("Üretim", 1, 40, D(20000)), ("Satış", 1, 5, D(15000)), ("Tüm", 1, 60, D(15000))])
        self.assertEqual([(a["ad"], a["egitim"], a["adam_saat"]) for a in r["aylik"] if a["egitim"]], [("Ocak", 1, D(240)), ("Şubat", 1, D(320)), ("Nisan", 1, D(200)), ("Temmuz", 1, D(128))])

    def test_butce_sinirlari(self):
        r = main.planla(IHT())
        self.assertEqual((len(r["plan"]), r["disi"], r["toplam"]["maliyet"], r["toplam"]["kalan"], r["uyarilar"]), (6, [], D(138000), None, []))
        r = main.planla(IHT(), D(60000))
        self.assertEqual(([x["egitim"] for x in r["plan"]], [x["egitim"] for x in r["disi"]], r["toplam"]["kalan"], sorted((u["onem"], u["kim"]) for u in r["uyarilar"])),
                         (["Yangın", "İSG", "İngilizce"], ["Liderlik", "Excel", "Sunum"], D(10000), [("Bilgi", "Excel · Finans"), ("Bilgi", "Sunum · Satış"), ("Orta", "Liderlik · Yönetim")]))
        r = main.planla(IHT(), D(30000))
        self.assertEqual(([x["egitim"] for x in r["plan"]], r["toplam"]["kalan"], len(r["disi"]), [u["tur"] for u in r["uyarilar"] if u["onem"] == "Yüksek"]), (["Yangın", "İSG"], D(-5000), 4, ["Zorunlu eğitimler bütçeyi aşıyor"]))
        self.assertEqual([x["egitim"] for x in main.planla(IHT(), D(35000))["plan"]], ["Yangın", "İSG"])

    def test_ay_atamasi(self):
        r = main.planla(IHT(), D(100000), {7, 2})
        self.assertEqual((ozet(r)[0][:2], ozet(r)[1][:2], [x[:2] for x in ozet(r)][2:], [u["tur"] for u in r["uyarilar"] if u["onem"] == "Orta"]), (("Yangın", 1), ("İSG", 2), [("Liderlik", 4), ("İngilizce", 10)], ["Tercih edilen ay kapalı"]))
        r = main.planla([e(f"E{i:02d}", "D", 1, 1) for i in range(1, 15)])
        self.assertEqual([x["ay"] for x in sorted(r["plan"], key=lambda x: int(x["egitim"][1:]))], [1, 7, 4, 10, 2, 8, 5, 11, 3, 9, 6, 12, 1, 7])
        r = main.planla([e("Büyük", "D", 10, 10, ay=1), e("K1", "D", 1, 1), e("K2", "D", 1, 1)], kapali=set(range(2, 12)))
        self.assertEqual([(x["egitim"], x["ay"]) for x in r["plan"]], [("Büyük", 1), ("K1", 12), ("K2", 12)])
        for iht, kap in (([], None), (IHT(), set(range(1, 13)))):
            with self.assertRaises(ValueError):
                main.planla(iht, None, kap)
        self.assertEqual([main.ay_coz(x) for x in ("3", "Mart", "ŞUBAT", "13", "0", "x")], [3, 3, 2, None, None, None])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "e.xlsx"
            r = main.calistir(ORNEK / "egitim_ihtiyaclari.csv", cikti, 2027, D(100000))
            self.assertEqual([(x["egitim"], x["ay"], x["maliyet"]) for x in r["plan"]],
                             [("Yangın ve Acil Durum", 1, D("15000.00")), ("Temel İş Sağlığı ve Güvenliği", 2, D("20000.00")), ("İş İngilizcesi", 4, D("15000.00")), ("Liderlik Gelişimi", 7, D("50000.00"))])
            self.assertEqual((r["toplam"]["maliyet"], r["toplam"]["disi"], len(r["uyarilar"])), (D("100000.00"), 2, 2))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Plan", "Aylık Özet", "Departman Özeti", "Bütçe Dışı", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "i.csv"
            y.write_text("Konu;Birim;Kişi Sayısı;Saat;Birim Maliyet;Sabit;Öncelik;Zorunlu;Ay\nA;;10;8;1.250,50;;;x;3\nB;D;0;8;;;;;\nC;D;2,5;8;;;;;\nD;D;5;8;;;0;;\nE;D;5;8;;;;;on üç\nF;D;5;;;;2;;Mart\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["egitim"], x["departman"], x["katilimci"], x["kisi"], x["oncelik"], x["zorunlu"], x["ay"], x["sure"]) for x in l], [u["tur"] for u in uy]),
                             ([("A", "Genel", 10, D("1250.50"), 3, True, 3, D(8)), ("F", "D", 5, D(0), 2, False, 3, D(0))], ["Eğitim satırı okunamadı"] * 4))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--ihtiyaclar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--butce", "50.000", "--yil", "2027", "--kapali-aylar", "7, Ağustos"]), 0)
            self.assertEqual(main.main(arg), 0)
            for ek in (["--butce", "x"], ["--kapali-aylar", "13"], ["--yil", "27"]):
                self.assertEqual(main.main(arg + ek), 2)
            self.assertEqual(main.main(arg + ["--kapali-aylar", "1,2,3,4,5,6,7,8,9,10,11,12"]), 1)
            self.assertEqual(main.main(["--ihtiyaclar", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Eğitim;Departman\nA;B\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
