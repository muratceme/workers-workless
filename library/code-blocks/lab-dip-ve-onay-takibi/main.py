"""
Lab Dip ve Onay Takibi — Workers / Workless kod bloğu
Müşteri Temsilciliği (Merchandising) › Production Merchandiser

Renk (lab dip), baskı ve aksesuar onaylarının gönderim ve müşteri cevabı durumunu sipariş bazında izler.
  Onay kalemi           = sipariş + onay türü + renk/detay;   her satır bir gönderim turudur, kalemin durumunu son tur belirler
  Gönderilmedi          = son turda gönderim tarihi yok
  Cevap bekleniyor      = gönderilmiş, cevap yok;   bekleme --cevap (7) günü aşmışsa "Cevap gecikti"
  Onaylandı / Reddedildi / Yorumlu = müşteri cevabının sonucu;   ret ve yorumda yeniden gönderim beklenir
  Termin geçti          = onaylanmamış kalemde onay termini < durum tarihi;   yaklaşıyor = termine --gun (7) gün veya az
  Onay oranı %          = onaylanan kalem ÷ toplam kalem × 100
  İlk turda onay %      = tek turda onaylanan kalem ÷ onaylanan kalem × 100
  Cevap süresi (gün)    = cevap tarihi − gönderim tarihi
Onay terminini siz yazarsınız (ör. kumaş ya da aksesuar siparişinin verilmesi gereken tarih); kod üretim planını bilmez.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"siparis": ("siparis no", "order no", "po no", "siparis"), "musteri": ("musteri", "buyer", "alici"), "model": ("model", "model no", "stil", "style"), "tur": ("onay turu", "tur", "onay tipi"),
         "detay": ("renk detay", "renk", "detay", "kalem"), "tur_no": ("tur", "gonderim turu", "tur no", "kacinci gonderim"), "gonderim": ("gonderim tarihi", "gonderilme tarihi"),
         "cevap": ("cevap tarihi", "musteri cevap tarihi", "yanit tarihi"), "sonuc": ("sonuc", "musteri karari", "durum"), "termin": ("onay termini", "termin", "son onay tarihi"), "sorumlu": ("sorumlu", "merchandiser")}
SONUC = {"onay": "Onay", "onaylandi": "Onay", "ok": "Onay", "approved": "Onay", "kabul": "Onay", "ret": "Ret", "red": "Ret", "reddedildi": "Ret", "rejected": "Ret", "yorum": "Yorum", "yorumlu": "Yorum", "revize": "Yorum",
         "sartli": "Yorum", "sartli onay": "Yorum"}
DURUM = ("Reddedildi, yeniden gönderilecek", "Yorumlu, yeniden gönderilecek", "Gönderilmedi", "Cevap gecikti", "Sonuç yazılmamış", "Cevap bekleniyor", "Onaylandı")
DURUM_RENK = {"Reddedildi, yeniden gönderilecek": KIRMIZI, "Yorumlu, yeniden gönderilecek": SARI, "Gönderilmedi": SARI, "Cevap gecikti": SARI, "Sonuç yazılmamış": SARI, "Cevap bekleniyor": MAVI, "Onaylandı": YESIL}
TERMIN = ("Termin geçti", "Termin yaklaşıyor", "")


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("siparis", "tur")):
        siparis = metin(r["siparis"])
        if not siparis:
            continue
        kim = f"{siparis} · {metin(r['tur'])} · {metin(r['detay'])}"
        tur_no = para(r["tur_no"]) if metin(r["tur_no"]) else None
        if not metin(r["tur"]) or (metin(r["tur_no"]) and (tur_no is None or tur_no < 1)):
            uy.append(uyari("Yüksek", "Satır okunamadı", kim, f"Satır {r['_satir']}: onay türü boş ya da tur numarası '{metin(r['tur_no'])}' okunamadı. Satır alınmadı."))
            continue
        for alan, ad in (("gonderim", "Gönderim tarihi"), ("cevap", "Cevap tarihi"), ("termin", "Onay termini")):
            if metin(r[alan]) and tarih(r[alan]) is None:
                uy.append(uyari("Orta", "Tarih okunamadı", kim, f"{ad} '{metin(r[alan])}' okunamadı; boş sayıldı."))
        sonuc = SONUC.get(katla(r["sonuc"]))
        if sonuc is None and metin(r["sonuc"]):
            uy.append(uyari("Orta", "Sonuç tanınmadı", kim, f"'{metin(r['sonuc'])}' Onay, Ret ya da Yorum olarak okunamadı; sonuç yazılmamış sayıldı."))
        l.append({"siparis": siparis, "musteri": metin(r["musteri"]) or "(müşteri yok)", "model": metin(r["model"]), "tur": metin(r["tur"]), "detay": metin(r["detay"]), "tur_no": int(tur_no) if tur_no is not None else None,
                  "gonderim": tarih(r["gonderim"]), "cevap": tarih(r["cevap"]), "sonuc": sonuc, "termin": tarih(r["termin"]), "sorumlu": metin(r["sorumlu"]), "satir": r["_satir"]})
    return l, uy


def _oran(pay: int, payda: int) -> Decimal | None:
    return yuvarla(Decimal(pay * 100) / payda) if payda else None


def _ort(d: list[int]) -> Decimal | None:
    return yuvarla(Decimal(sum(d)) / len(d), "0.1") if d else None


def degerlendir(satirlar: list[dict], bugun: date, cevap: int = 7, gun: int = 7, cok_tur: int = 3) -> dict:
    if not satirlar:
        raise ValueError("İzlenecek onay kaydı yok.")
    grup, uy = {}, []
    for s in satirlar:
        grup.setdefault((katla(s["siparis"]), katla(s["tur"]), katla(s["detay"])), []).append(s)
    kalemler = []
    for g in grup.values():
        g.sort(key=lambda s: (s["tur_no"] if s["tur_no"] is not None else 0, s["gonderim"] or date.max, s["satir"]))
        son = g[-1]
        k = {a: son[a] for a in ("siparis", "musteri", "model", "tur", "detay", "gonderim", "cevap", "sorumlu")}
        k["termin"] = son["termin"] or next((s["termin"] for s in reversed(g) if s["termin"]), None)
        k["tur_sayisi"], k["ret"] = sum(1 for s in g if s["gonderim"]), sum(1 for s in g if s["sonuc"] in ("Ret", "Yorum"))
        k["sureler"] = [(s["cevap"] - s["gonderim"]).days for s in g if s["gonderim"] and s["cevap"] and s["cevap"] >= s["gonderim"]]
        k["bekleme"] = None
        if son["gonderim"] is None:
            k["durum"] = "Gönderilmedi"
        elif son["cevap"] is None:
            k["bekleme"] = max((bugun - son["gonderim"]).days, 0)
            k["durum"] = "Cevap gecikti" if k["bekleme"] > cevap else "Cevap bekleniyor"
        else:
            k["durum"] = {"Onay": "Onaylandı", "Ret": "Reddedildi, yeniden gönderilecek", "Yorum": "Yorumlu, yeniden gönderilecek"}.get(son["sonuc"], "Sonuç yazılmamış")
        k["onayli"] = k["durum"] == "Onaylandı"
        k["termin_kalan"] = (k["termin"] - bugun).days if k["termin"] and not k["onayli"] else None
        k["termin_durum"] = "" if k["termin_kalan"] is None or k["termin_kalan"] > gun else "Termin geçti" if k["termin_kalan"] < 0 else "Termin yaklaşıyor"
        kalemler.append(k)
        kim = f"{k['siparis']} · {k['tur']} · {k['detay']}" if k["detay"] else f"{k['siparis']} · {k['tur']}"
        if k["termin_durum"] == "Termin geçti":
            uy.append(uyari("Yüksek", "Onay termini geçti", kim, f"Termin {k['termin']:%d.%m.%Y} ({-k['termin_kalan']} gün önce); durum: {k['durum']}. {k['musteri']}."))
        elif k["termin_durum"]:
            uy.append(uyari("Orta", "Onay termini yaklaşıyor", kim, f"Termin {k['termin']:%d.%m.%Y} ({k['termin_kalan']} gün kaldı); durum: {k['durum']}. {k['musteri']}."))
        if k["durum"] == "Cevap gecikti":
            uy.append(uyari("Orta", "Müşteri cevabı gecikti", kim, f"{k['gonderim']:%d.%m.%Y} tarihinde gönderildi, {k['bekleme']} gündür cevap yok. {k['musteri']}."))
        elif k["durum"] in DURUM[:2]:
            uy.append(uyari("Orta", "Yeniden gönderim bekliyor", kim, f"{k['cevap']:%d.%m.%Y} tarihli cevap: {k['durum'].split(',')[0].lower()}. {k['tur_sayisi']}. turdan sonra yeni gönderim yapılmadı."))
        elif k["durum"] == "Sonuç yazılmamış":
            uy.append(uyari("Orta", "Cevap sonucu yazılmamış", kim, f"Cevap tarihi {k['cevap']:%d.%m.%Y} ama sonuç (Onay / Ret / Yorum) yok."))
        if k["tur_sayisi"] >= cok_tur and not k["onayli"]:
            uy.append(uyari("Orta", "Çok turlu onay", kim, f"{k['tur_sayisi']} kez gönderildi, hâlâ onay yok. Müşteriyle beklentiyi netleştirin."))

    def ozet(l: list[dict]) -> dict:
        onayli, sureler = [k for k in l if k["onayli"]], [x for k in l for x in k["sureler"]]
        acik_termin = [k["termin"] for k in l if not k["onayli"] and k["termin"]]
        return {"kalem": len(l), "onayli": len(onayli), "acik": len(l) - len(onayli), "oran": _oran(len(onayli), len(l)), "ilk_tur": _oran(sum(1 for k in onayli if k["tur_sayisi"] <= 1), len(onayli)),
                "gecmis": sum(1 for k in l if k["termin_durum"] == "Termin geçti"), "yakin": sum(1 for k in l if k["termin_durum"] == "Termin yaklaşıyor"), "bekleyen": sum(1 for k in l if k["durum"] in ("Cevap bekleniyor", "Cevap gecikti")),
                "ort_cevap": _ort(sureler), "en_yakin": min(acik_termin, default=None)}

    def kirilim(alan: str) -> list[dict]:
        g2 = {}
        for k in kalemler:
            g2.setdefault(katla(k[alan]), []).append(k)
        return sorted(({"ad": g[0][alan], "musteri": g[0]["musteri"], "model": g[0]["model"], **ozet(g)} for g in g2.values()), key=lambda s: (-s["gecmis"], -s["acik"], katla(s["ad"])))

    kalemler.sort(key=lambda k: (k["onayli"], TERMIN.index(k["termin_durum"]), DURUM.index(k["durum"]), k["termin"] or date.max, katla(k["siparis"]), katla(k["tur"]), katla(k["detay"])))
    return {"kalemler": kalemler, "ozet": ozet(kalemler), "siparisler": kirilim("siparis"), "musteriler": kirilim("musteri"), "turler": kirilim("tur"),
            "durumlar": [{"durum": d, "kalem": sum(1 for k in kalemler if k["durum"] == d)} for d in DURUM], "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, bugun: date) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Onay Kalemleri"
    _baslik(ws, ["Durum", "Termin Durumu", "Sipariş No", "Müşteri", "Model", "Onay Türü", "Renk / Detay", "Tur Sayısı", "Son Gönderim", "Son Cevap", "Cevap Bekleme (gün)", "Onay Termini", "Termine Kalan (gün)", "Sorumlu", "Not (siz doldurun)"],
            (30, 17, 13, 22, 14, 14, 20, 8, 12, 12, 10, 12, 10, 18, 30))
    for k in r["kalemler"]:
        _satir(ws, [k["durum"], k["termin_durum"] or None, k["siparis"], k["musteri"], k["model"], k["tur"], k["detay"], k["tur_sayisi"] or None, k["gonderim"], k["cevap"], k["bekleme"], k["termin"], k["termin_kalan"], k["sorumlu"], None],
               tarih_s=(9, 10, 12), dolgu={1: DURUM_RENK[k["durum"]], 2: KIRMIZI if k["termin_durum"] == "Termin geçti" else SARI if k["termin_durum"] else None})
    ws = wb.create_sheet("Özet")
    o = r["ozet"]
    _baslik(ws, ["Gösterge", "Değer"], (44, 14))
    for etiket, deger, para_mi in (("Onay kalemi", o["kalem"], False), ("Onaylanan", o["onayli"], False), ("Açık kalem", o["acik"], False), ("Onay oranı %", o["oran"], True), ("İlk turda onay %", o["ilk_tur"], True),
                                   ("Müşteri cevabı beklenen", o["bekleyen"], False), ("Termini geçmiş kalem", o["gecmis"], False), ("Termini yaklaşan kalem", o["yakin"], False), ("Ortalama müşteri cevap süresi (gün)", o["ort_cevap"], True)):
        _satir(ws, [etiket, deger], para_s=(2,) if para_mi else ())
    ws.append([])
    _baslik(ws, ["Durum", "Kalem"], dondur="A2")
    for d in r["durumlar"]:
        _satir(ws, [d["durum"], d["kalem"]], dolgu={1: DURUM_RENK[d["durum"]] if d["kalem"] else None})
    ws.append([])
    ws.append([f"Durum tarihi: {bugun:%d.%m.%Y}"])
    for baslik, alan, adlar in (("Siparişler", "siparisler", ["Sipariş No", "Müşteri", "Model"]), ("Müşteriler", "musteriler", ["Müşteri"]), ("Onay Türleri", "turler", ["Onay Türü"])):
        ws = wb.create_sheet(baslik)
        n = len(adlar)
        _baslik(ws, adlar + ["Kalem", "Onaylanan", "Açık", "Onay %", "İlk Turda Onay %", "Cevap Beklenen", "Termini Geçmiş", "Termini Yaklaşan", "Ort. Cevap Süresi (gün)", "En Yakın Açık Termin"], (22,) * n + (8, 10, 8, 9, 11, 10, 10, 10, 11, 13))
        for s in r[alan]:
            _satir(ws, [s["ad"]] + ([s["musteri"], s["model"]] if n == 3 else []) + [s["kalem"], s["onayli"], s["acik"] or None, s["oran"], s["ilk_tur"], s["bekleyen"] or None, s["gecmis"] or None, s["yakin"] or None, s["ort_cevap"], s["en_yakin"]],
                   para_s=(n + 4, n + 5), tarih_s=(n + 10,), dolgu={n + 7: KIRMIZI if s["gecmis"] else None})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, cevap: int = 7, gun: int = 7, cok_tur: int = 3) -> dict:
    satirlar, uy = oku(yol)
    r = degerlendir(satirlar, bugun, cevap, gun, cok_tur)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, bugun)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Renk (lab dip), baskı ve aksesuar onaylarının gönderim ve onay durumunu sipariş bazında takip eder.")
    p.add_argument("--onaylar", type=Path, help="Sipariş No, Müşteri, Model, Onay Türü, Renk / Detay, Tur, Gönderim Tarihi, Cevap Tarihi, Sonuç, Onay Termini, Sorumlu")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--cevap", type=int, default=7, help="Gönderimden sonra kaç günden fazla cevapsızlık gecikme sayılsın (varsayılan 7)")
    p.add_argument("--gun", type=int, default=7, help="Onay terminine kaç gün kala uyarılsın (varsayılan 7)")
    p.add_argument("--tur", type=int, default=3, help="Bu kadar ve üstü gönderimde hâlâ onaylanmayan kalem uyarılır (varsayılan 3)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "lab_dip_onay_takibi.xlsx")
    a = p.parse_args(argv)
    ornek = a.onaylar is None
    yol = ORNEK / "onay_kayitlari.csv" if ornek else a.onaylar
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    if bugun is None or a.cevap < 0 or a.gun < 0 or a.tur < 2:
        print("[X] --tarih GG.AA.YYYY; --cevap ve --gun eksi olamaz; --tur en az 2 olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, bugun, a.cevap, a.gun, a.tur)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {o['kalem']} onay kalemi · onaylanan {o['onayli']} (%{o['oran']}) · açık {o['acik']} · cevap beklenen {o['bekleyen']} · termini geçmiş {o['gecmis']}, yaklaşan {o['yakin']}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
