import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)
_n = [0]


def g(n):
    return None if n is None else BUGUN + timedelta(days=n)


def s(detay, gonderim=None, cevap=None, sonuc=None, termin=30, tur_no=1, siparis="S1", tur="Lab dip", musteri="M"):
    """Tarihler: bugünden kaç gün sonra (eksi = geçmiş; None = boş)."""
    _n[0] += 1
    return {"siparis": siparis, "musteri": musteri, "model": "", "tur": tur, "detay": detay, "tur_no": tur_no, "gonderim": g(gonderim), "cevap": g(cevap), "sonuc": sonuc, "termin": g(termin), "sorumlu": "", "satir": _n[0]}


class DegerlendirTesti(unittest.TestCase):
    def test_durumlar(self):
        l = [s("onaylı", -40, -33, "Ret", -5), s("ONAYLI", -27, -22, "Onay", -5, 2), s("bekleyen7", -7, termin=3), s("geciken8", -8, termin=-2), s("gönderilmemiş", termin=8), s("yorumlu", -11, -6, "Yorum", 7),
             s("sonuçsuz", -20, -15), s("ret", -10, -3, "Ret", None), s("ilk_turda", -30, -26, "Onay", -1, tur="Baskı", siparis="S2", musteri="N")]
        r = main.degerlendir(l, BUGUN)
        self.assertEqual({k["detay"]: (k["durum"], k["termin_durum"], k["tur_sayisi"], k["bekleme"], k["termin_kalan"]) for k in r["kalemler"]},
                         {"ONAYLI": ("Onaylandı", "", 2, None, None), "bekleyen7": ("Cevap bekleniyor", "Termin yaklaşıyor", 1, 7, 3), "geciken8": ("Cevap gecikti", "Termin geçti", 1, 8, -2), "gönderilmemiş": ("Gönderilmedi", "", 0, None, 8),
                          "yorumlu": ("Yorumlu, yeniden gönderilecek", "Termin yaklaşıyor", 1, None, 7), "sonuçsuz": ("Sonuç yazılmamış", "", 1, None, 30), "ret": ("Reddedildi, yeniden gönderilecek", "", 1, None, None),
                          "ilk_turda": ("Onaylandı", "", 1, None, None)})
        self.assertEqual([k["detay"] for k in r["kalemler"]], ["geciken8", "yorumlu", "bekleyen7", "ret", "gönderilmemiş", "sonuçsuz", "ONAYLI", "ilk_turda"])
        self.assertEqual(r["ozet"], {"kalem": 8, "onayli": 2, "acik": 6, "oran": D("25.00"), "ilk_tur": D("50.00"), "gecmis": 1, "yakin": 2, "bekleyen": 2, "ort_cevap": D("5.5"), "en_yakin": g(-2)})
        self.assertEqual([(x["ad"], x["kalem"], x["onayli"], x["acik"], x["gecmis"], x["ort_cevap"]) for x in r["siparisler"]], [("S1", 7, 1, 6, 1, D("5.8")), ("S2", 1, 1, 0, 0, D("4.0"))])
        self.assertEqual(([(x["ad"], x["kalem"]) for x in r["turler"]], [(x["ad"], x["kalem"]) for x in r["musteriler"]]), ([("Lab dip", 7), ("Baskı", 1)], [("M", 7), ("N", 1)]))
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         sorted([("Yüksek", "Onay termini geçti", "S1 · Lab dip · geciken8"), ("Orta", "Müşteri cevabı gecikti", "S1 · Lab dip · geciken8"), ("Orta", "Onay termini yaklaşıyor", "S1 · Lab dip · bekleyen7"),
                                 ("Orta", "Onay termini yaklaşıyor", "S1 · Lab dip · yorumlu"), ("Orta", "Yeniden gönderim bekliyor", "S1 · Lab dip · yorumlu"), ("Orta", "Yeniden gönderim bekliyor", "S1 · Lab dip · ret"),
                                 ("Orta", "Cevap sonucu yazılmamış", "S1 · Lab dip · sonuçsuz")]))

    def test_esikler_ve_turlar(self):
        d = lambda **kw: main.degerlendir([s("a", -5, termin=10)], BUGUN, **kw)["kalemler"][0]   # noqa: E731
        self.assertEqual(((d()["durum"], d()["termin_durum"]), (d(cevap=4)["durum"], d(gun=10)["termin_durum"])), (("Cevap bekleniyor", ""), ("Cevap gecikti", "Termin yaklaşıyor")))
        l = [s("x", -30, -25, "Ret", tur_no=None), s("x", -20, -15, "Yorum", tur_no=None), s("x", -10, None, None, None, tur_no=None)]
        r = main.degerlendir(l, BUGUN)
        k = r["kalemler"][0]
        self.assertEqual((k["durum"], k["tur_sayisi"], k["ret"], k["bekleme"], k["termin"], sorted(u["tur"] for u in r["uyarilar"])), ("Cevap gecikti", 3, 2, 10, g(30), ["Müşteri cevabı gecikti", "Çok turlu onay"]))
        self.assertEqual([u["tur"] for u in main.degerlendir(l, BUGUN, cok_tur=4, cevap=10)["uyarilar"]], [])
        with self.assertRaises(ValueError):
            main.degerlendir([], BUGUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "onay_kayitlari.csv", cikti, BUGUN)
            self.assertEqual([(k["siparis"], k["detay"], k["durum"], k["termin_durum"], k["bekleme"]) for k in r["kalemler"]],
                             [("PO-5001", "Göğüs logosu", "Cevap gecikti", "Termin geçti", 11), ("PO-5001", "Kırmızı", "Cevap bekleniyor", "Termin yaklaşıyor", 7), ("PO-5002", "Bej", "Yorumlu, yeniden gönderilecek", "", None),
                              ("PO-5002", "Düğme", "Gönderilmedi", "", None), ("PO-5002", "Siyah", "Sonuç yazılmamış", "", None), ("PO-5001", "Fermuar", "Onaylandı", "", None), ("PO-5001", "Lacivert", "Onaylandı", "", None),
                              ("PO-5002", "Yaka nakışı", "Onaylandı", "", None)])
            self.assertEqual(r["ozet"], {"kalem": 8, "onayli": 3, "acik": 5, "oran": D("37.50"), "ilk_tur": D("66.67"), "gecmis": 1, "yakin": 1, "bekleyen": 2, "ort_cevap": D("5.3"), "en_yakin": date(2026, 10, 10)})
            self.assertEqual([(x["ad"], x["kalem"], x["onayli"], x["acik"], x["gecmis"], x["en_yakin"]) for x in r["siparisler"]], [("PO-5001", 4, 2, 2, 1, date(2026, 10, 10)), ("PO-5002", 4, 1, 3, 0, date(2026, 10, 25))])
            self.assertEqual([(u["onem"], u["tur"], u["kim"].split(" · ")[-1]) for u in r["uyarilar"]],
                             [("Yüksek", "Onay termini geçti", "Göğüs logosu"), ("Orta", "Cevap sonucu yazılmamış", "Siyah"), ("Orta", "Müşteri cevabı gecikti", "Göğüs logosu"), ("Orta", "Onay termini yaklaşıyor", "Kırmızı"),
                              ("Orta", "Yeniden gönderim bekliyor", "Bej")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Onay Kalemleri", "Özet", "Siparişler", "Müşteriler", "Onay Türleri", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "o.csv"
            y.write_text("Order No;Buyer;Style;Tür;Renk;Gönderilme Tarihi;Yanıt Tarihi;Müşteri Kararı;Termin\nP1;B;S;Lab dip;Mavi;01.10.2026;x;APPROVED;15.10.2026\nP1;;;Baskı;;;;Belki;\nP2;;;;;;;;\n;;;Lab dip;;;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["siparis"], x["musteri"], x["model"], x["tur"], x["detay"], x["tur_no"], x["gonderim"], x["cevap"], x["sonuc"], x["termin"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("P1", "B", "S", "Lab dip", "Mavi", None, date(2026, 10, 1), None, "Onay", date(2026, 10, 15)), ("P1", "(müşteri yok)", "", "Baskı", "", None, None, None, None, None)],
                              [("Tarih okunamadı", "P1 · Lab dip · Mavi"), ("Sonuç tanınmadı", "P1 · Baskı ·"), ("Satır okunamadı", "P2 ·  ·")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--onaylar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "12.10.2026", "--cevap", "5", "--gun", "10", "--tur", "2"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--cevap", "-1"], ["--gun", "-1"], ["--tur", "1"])], [2] * 4)
            self.assertEqual(main.main(["--onaylar", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Sipariş No;Müşteri\nA;B\n", "Sipariş No;Onay Türü\n;Lab dip\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
