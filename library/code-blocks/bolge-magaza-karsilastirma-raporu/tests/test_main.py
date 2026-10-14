import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def m(ad, bolge, alan, ciro, gy, fis, ziy, smm, stok):
    c = lambda x: None if x is None else D(x)   # noqa: E731
    return {"magaza": ad, "bolge": bolge, "alan": c(alan), "ciro": D(ciro), "gy": c(gy), "fis": c(fis), "ziyaretci": c(ziy), "smm": c(smm), "stok": c(stok)}


def veri():
    return [m("M1", "A", 100, 1000, 800, 10, 100, 600, 300), m("M2", "A", 100, 500, 500, 10, 200, 300, 300), m("M3", "a", None, 300, None, None, None, None, None)]


class KarsilastirTesti(unittest.TestCase):
    def test_gostergeler(self):
        r = main.karsilastir(veri())
        b = r["bolgeler"][0]
        self.assertEqual((len(r["bolgeler"]), b["bolge"], b["sayi"], b["m2"], b["lfl"], b["donusum"], b["sepet"], b["devir"], b["stok_gun"], r["toplam"]),
                         (1, "A", 3, D("7.50"), D("15.38"), D("6.67"), D("75.00"), D("1.50"), D("20.00"), None))
        self.assertEqual([(s["magaza"], s["ort_sira"], s["pay"], s["m2"], s["m2_endeks"], s["lfl"], s["donusum"], s["donusum_endeks"], s["devir"], s["devir_endeks"], s["stok_gun"]) for s in r["magazalar"]],
                         [("M1", D("1.00"), D("55.56"), D("10.00"), D("133.33"), D("25.00"), D("10.00"), D("149.93"), D("2.00"), D("133.33"), D("15.00")),
                          ("M2", D("2.00"), D("27.78"), D("5.00"), D("66.67"), D("0.00"), D("5.00"), D("74.96"), D("1.00"), D("66.67"), D("30.00")),
                          ("M3", None, D("16.67"), None, None, None, None, None, None, None, None)])
        self.assertEqual(sorted(u["tur"] for u in r["uyarilar"]), ["Dönüşüm düşük", "Metrekare verimliliği düşük", "Stok devri yavaş"])

    def test_esikler_ve_gun(self):
        tur = lambda **k: sorted(u["tur"] for u in main.karsilastir(veri(), **k)["uyarilar"])   # noqa: E731
        self.assertEqual(tur(sapma=D(30)), ["Metrekare verimliliği düşük", "Stok devri yavaş"])
        self.assertEqual(tur(lfl_esik=D(1), sapma=D(40)), ["Geçen yılın gerisinde"])
        self.assertEqual(main.karsilastir(veri(), gun=D(90))["bolgeler"][0]["stok_gun"], D("60.00"))
        iki = main.karsilastir(veri() + [m("X", "B", 50, 100, None, None, None, None, None)])
        self.assertEqual(([b["bolge"] for b in iki["bolgeler"]], iki["toplam"]["ciro"], iki["toplam"]["m2"]), (["A", "B"], D(1900), D("6.40")))
        with self.assertRaises(ValueError):
            main.karsilastir([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "magaza_donem_verileri.csv", cikti)
            self.assertEqual([(b["bolge"], b["ciro"], b["m2"], b["lfl"], b["donusum"], b["devir"]) for b in r["bolgeler"]],
                             [("Ege", D(3550000), D("2088.24"), D("2.90"), D("20.92"), D("1.33")), ("Marmara", D(5650000), D("2456.52"), D("4.29"), D("19.58"), D("1.23"))])
            self.assertEqual([(s["magaza"], s["ort_sira"]) for s in r["magazalar"]],
                             [("Alsancak", D("1.25")), ("Karşıyaka", D("1.25")), ("Bornova", D("3.00")), ("Merkez", D("1.00")), ("Kadıköy", D("2.00")), ("AVM", D("3.00"))])
            self.assertEqual([(u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Dönüşüm düşük", "Ege · Bornova"), ("Dönüşüm düşük", "Marmara · AVM"), ("Geçen yılın gerisinde", "Ege · Bornova"), ("Metrekare verimliliği düşük", "Ege · Bornova"),
                              ("Stok devri yavaş", "Ege · Bornova")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Mağazalar", "Bölgeler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "m.csv"
            y.write_text("Şube;Satış Tutarı;Alan\nA;1.000;10\na;5;5\nB;x;1\nC;5;-1\nD;5;\n;5;5\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["magaza"], x["bolge"], x["ciro"], x["alan"]) for x in l], [u["tur"] for u in uy]),
                             ([("A", "Bölge", D(1000), D(10)), ("D", "Bölge", D(5), None)], ["Mağaza iki kez yazılmış", "Satır okunamadı", "Satır okunamadı", "Satış alanı yok"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--magazalar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--gun", "91", "--sapma", "10", "--lfl", "0"]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--gun", "0"], ["--sapma", "100"], ["--lfl", "x"])], [2] * 3)
            self.assertEqual(main.main(["--magazalar", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Mağaza;Alan\nA;5\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
