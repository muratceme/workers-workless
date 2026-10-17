"""
Bölge Mağaza Karşılaştırma Raporu — Gerisi Bende kod bloğu
Mağaza Operasyonları › Bölge Müdürü
Bölgedeki mağazaları metrekare verimliliği, LFL büyüme, dönüşüm ve stok devir hızına göre karşılaştırır.
  Metrekare verimliliği = dönem cirosu ÷ satış alanı (m²)
  LFL büyüme %          = (ciro − geçen yıl ciro) ÷ geçen yıl ciro × 100
  Dönüşüm oranı %       = fiş adedi ÷ ziyaretçi sayısı × 100
  Sepet ortalaması      = ciro ÷ fiş adedi
  Stok devir hızı       = satılan malın maliyeti ÷ ortalama stok (maliyet bedeliyle) — dönem içinde
  Stok günü             = ortalama stok ÷ satılan malın maliyeti × dönem gün sayısı (--gun)
  Bölgeye göre endeks   = mağazanın değeri ÷ bölgenin değeri × 100 (bölge değeri toplamların oranıdır)
  Uyarı                 = endeks 100 − --sapma (20) altında · LFL --lfl (−%5) altında
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"magaza": ("magaza", "magaza adi", "sube"), "bolge": ("bolge", "bolge adi"), "alan": ("satis alani", "satis alani m2", "alan", "m2", "metrekare"),
         "ciro": ("ciro", "net satis", "satis tutari", "satis"), "gy": ("gecen yil ciro", "gecen yil", "onceki yil ciro"), "fis": ("fis adedi", "fis sayisi", "islem adedi"),
         "ziyaretci": ("ziyaretci", "ziyaretci sayisi", "trafik"), "smm": ("satilan malin maliyeti", "smm", "satis maliyeti"), "stok": ("ortalama stok", "ortalama stok maliyet", "stok")}
SAYI = ("alan", "ciro", "gy", "fis", "ziyaretci", "smm", "stok")
SIFIR = Decimal(0)
OLCU = (("m2", "Metrekare verimliliği düşük", "m² başına ciro"), ("donusum", "Dönüşüm düşük", "dönüşüm oranı"), ("devir", "Stok devri yavaş", "stok devir hızı"))


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("magaza", "ciro")):
        magaza = metin(r["magaza"])
        if not magaza:
            continue
        d = {a: para(r[a]) if metin(r[a]) else None for a in SAYI}
        if d["ciro"] is None or any(metin(r[a]) and (d[a] is None or d[a] < 0) for a in SAYI):
            uy.append(uyari("Yüksek", "Satır okunamadı", magaza, f"Satır {r['_satir']}: sayılardan biri okunamadı ya da eksi. Mağaza alınmadı."))
            continue
        bolge = metin(r["bolge"]) or "Bölge"
        anahtar = (katla(bolge), katla(magaza))
        if anahtar in gorulen:
            uy.append(uyari("Yüksek", "Mağaza iki kez yazılmış", magaza, f"Satır {r['_satir']}: aynı bölgede aynı mağaza ikinci kez geçiyor. İkinci satır alınmadı."))
            continue
        gorulen.add(anahtar)
        if not d["alan"]:
            uy.append(uyari("Orta", "Satış alanı yok", magaza, "Metrekare verimliliği hesaplanamadı."))
        l.append({"magaza": magaza, "bolge": bolge, **d})
    return l, uy


def _oran(pay, payda, kat=100) -> Decimal | None:
    return yuvarla(Decimal(pay) * kat / Decimal(payda)) if payda else None


def gosterge(l: list[dict], gun: Decimal) -> dict:
    def top(alan, *kosul):
        return sum((m[alan] for m in l if m[alan] is not None and all(m[k] for k in kosul)), SIFIR)
    ciro, fis = top("ciro"), top("fis")
    gy, smm, stok = top("gy", "gy"), top("smm", "stok"), top("stok", "smm")
    return {"sayi": len(l), "alan": top("alan"), "ciro": ciro, "gy": gy, "fis": fis, "ziyaretci": top("ziyaretci"), "smm": top("smm"), "stok": top("stok"),
            "m2": _oran(top("ciro", "alan"), top("alan"), 1), "lfl": _oran(top("ciro", "gy") - gy, gy), "donusum": _oran(top("fis", "ziyaretci"), top("ziyaretci", "fis")),
            "sepet": _oran(top("ciro", "fis"), fis, 1), "devir": _oran(smm, stok, 1), "stok_gun": _oran(stok * gun, smm, 1)}


def karsilastir(magazalar: list[dict], gun: Decimal = Decimal(30), sapma: Decimal = Decimal(20), lfl_esik: Decimal = Decimal(-5)) -> dict:
    if not magazalar:
        raise ValueError("Karşılaştırılacak mağaza yok.")
    gruplar, uy, sonuc, bolgeler = {}, [], [], []
    for m in magazalar:
        gruplar.setdefault(katla(m["bolge"]), []).append(m)
    alt = 100 - sapma
    for l in gruplar.values():
        b = {"bolge": l[0]["bolge"], **gosterge(l, gun)}
        bolgeler.append(b)
        satirlar = [{"magaza": m["magaza"], "bolge": b["bolge"], **gosterge([m], gun)} for m in l]
        for s in satirlar:
            for k in ("m2", "donusum", "devir"):
                s[k + "_endeks"] = _oran(s[k], b[k]) if s[k] is not None else None
        for k in ("m2", "lfl", "donusum", "devir"):
            degerli = sorted((s for s in satirlar if s[k] is not None), key=lambda s: -s[k])
            for s in satirlar:
                s[k + "_sira"] = next((i for i, x in enumerate(degerli, 1) if x[k] == s[k]), None) if s[k] is not None else None
        for s in satirlar:
            siralar = [s[k + "_sira"] for k in ("m2", "lfl", "donusum", "devir") if s[k + "_sira"] is not None]
            s["ort_sira"] = yuvarla(Decimal(sum(siralar)) / len(siralar)) if siralar else None
            s["pay"] = _oran(s["ciro"], b["ciro"])
            for k, tur, ad in OLCU:
                if s[k + "_endeks"] is not None and s[k + "_endeks"] < alt:
                    uy.append(uyari("Orta", tur, f"{s['bolge']} · {s['magaza']}", f"{ad.capitalize()} {s[k]:,.2f}, bölge {b[k]:,.2f} (endeks {s[k + '_endeks']}; eşik {alt})."))
            if s["lfl"] is not None and s["lfl"] < lfl_esik:
                uy.append(uyari("Orta", "Geçen yılın gerisinde", f"{s['bolge']} · {s['magaza']}", f"Ciro {s['ciro']:,.2f}, geçen yıl {s['gy']:,.2f} (LFL %{s['lfl']}; eşik %{lfl_esik})."))
        sonuc += sorted(satirlar, key=lambda s: (s["ort_sira"] is None, s["ort_sira"] or 0, -s["ciro"], katla(s["magaza"])))
    bolgeler.sort(key=lambda b: katla(b["bolge"]))
    sonuc.sort(key=lambda s: katla(s["bolge"]))
    return {"magazalar": sonuc, "bolgeler": bolgeler, "toplam": gosterge(magazalar, gun) if len(bolgeler) > 1 else None, "uyarilar": uy}


BASLIK = ["Ciro", "Satış Alanı m²", "m² Başına Ciro", "Geçen Yıl Ciro", "LFL Büyüme %", "Fiş Adedi", "Ziyaretçi", "Dönüşüm %", "Sepet Ortalaması", "SMM", "Ortalama Stok", "Stok Devir Hızı", "Stok Günü"]


def _deger(g: dict) -> list:
    return [g["ciro"], g["alan"] or None, g["m2"], g["gy"] or None, g["lfl"], g["fis"] or None, g["ziyaretci"] or None, g["donusum"], g["sepet"], g["smm"] or None, g["stok"] or None, g["devir"], g["stok_gun"]]


def rapor_yaz(cikti: Path, r: dict, gun: Decimal, sapma: Decimal, lfl_esik: Decimal) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Mağazalar"
    alt = 100 - sapma
    renk = lambda e: None if e is None else KIRMIZI if e < alt else YESIL if e >= 100 else None   # noqa: E731
    _baslik(ws, ["Bölge", "Mağaza", "Ortalama Sıra", "Ciro Payı %"] + BASLIK + ["m² Endeksi", "Dönüşüm Endeksi", "Stok Devir Endeksi", "m² Sırası", "LFL Sırası", "Dönüşüm Sırası", "Devir Sırası"],
            (14, 22, 9, 9, 15, 10, 12, 15, 10, 11, 11, 10, 11, 15, 15, 10, 10, 10, 10, 10, 8, 8, 8, 8), "C2")
    for s in r["magazalar"]:
        _satir(ws, [s["bolge"], s["magaza"], s["ort_sira"], s["pay"]] + _deger(s) + [s["m2_endeks"], s["donusum_endeks"], s["devir_endeks"], s["m2_sira"], s["lfl_sira"], s["donusum_sira"], s["devir_sira"]],
               para_s=(5, 7, 8, 13, 14, 15), dolgu={9: None if s["lfl"] is None else KIRMIZI if s["lfl"] < lfl_esik else YESIL if s["lfl"] >= 0 else None,
                                                  18: renk(s["m2_endeks"]), 19: renk(s["donusum_endeks"]), 20: renk(s["devir_endeks"])})
    ws.append([])
    ws.append([f"Endeks: mağaza ÷ bölge × 100; {alt} altı kırmızı. Ortalama sıra: dört göstergedeki sıraların ortalaması (1 en iyi). Stok günü {gun} günlük döneme göre."])
    ws = wb.create_sheet("Bölgeler")
    _baslik(ws, ["Bölge", "Mağaza Sayısı"] + BASLIK, (16, 9, 15, 10, 12, 15, 10, 11, 11, 10, 11, 15, 15, 10, 10))
    for b in r["bolgeler"] + ([{"bolge": "TOPLAM", **r["toplam"]}] if r["toplam"] else []):
        k = _satir(ws, [b["bolge"], b["sayi"]] + _deger(b), para_s=(3, 5, 6, 11, 12, 13))
        if b["bolge"] == "TOPLAM":
            for h in ws[k]:
                h.font = Font(bold=True)
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, gun: Decimal = Decimal(30), sapma: Decimal = Decimal(20), lfl_esik: Decimal = Decimal(-5)) -> dict:
    magazalar, uy = oku(yol)
    r = karsilastir(magazalar, gun, sapma, lfl_esik)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, gun, sapma, lfl_esik)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Bölgedeki mağazaları metrekare verimliliği, LFL büyüme, dönüşüm ve stok devir hızına göre karşılaştırır.")
    p.add_argument("--magazalar", type=Path, help="Mağaza, Bölge, Satış Alanı, Ciro, Geçen Yıl Ciro, Fiş Adedi, Ziyaretçi, Satılan Malın Maliyeti, Ortalama Stok")
    p.add_argument("--gun", default="30", help="Dönemin gün sayısı; stok günü için (varsayılan 30)")
    p.add_argument("--sapma", default="20", help="Bölgenin bu yüzde altında kalan gösterge uyarılır (varsayılan 20)")
    p.add_argument("--lfl", default="-5", help="LFL büyümesi bu yüzdenin altındaysa uyarılır (varsayılan -5)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "bolge_magaza_karsilastirma.xlsx")
    a = p.parse_args(argv)
    yol = a.magazalar or ORNEK / "magaza_donem_verileri.csv"
    gun, sapma, lfl = para(a.gun), para(a.sapma), para(a.lfl)
    if gun is None or gun <= 0 or sapma is None or not 0 <= sapma < 100 or lfl is None:
        print("[X] --gun sıfırdan büyük, --sapma 0-100 arası, --lfl bir yüzde olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, gun, sapma, lfl)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(r['bolgeler'])} bölge · {len(r['magazalar'])} mağaza")
    for b in r["bolgeler"]:
        print(f"     {b['bolge']}: ciro {b['ciro']:,.2f}" + (f" · m² {b['m2']:,.2f}" if b["m2"] is not None else "") + (f" · LFL %{b['lfl']}" if b["lfl"] is not None else "")
              + (f" · dönüşüm %{b['donusum']}" if b["donusum"] is not None else "") + (f" · devir {b['devir']}" if b["devir"] is not None else ""))
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
