"""
Stok ABC-XYZ Analizi — Workers / Workless kod bloğu
Üretim › Üretim Planlama / Malzeme Planlama

Stokları iki eksende sınıflandırır:
  - ABC (değer): yıllık tüketim değerine göre Pareto; varsayılan A ≤ %80, B ≤ %95, C kalan (kümülatif pay).
  - XYZ (talep değişkenliği): dönemsel tüketimin değişim katsayısı CV = standart sapma / ortalama;
    varsayılan X ≤ 0,5, Y ≤ 1,0, Z > 1,0.
Dokuz sınıfın her biri için stok politikası önerir; hiç tüketimi olmayan kalemleri "Hareketsiz", geçmişi kısa
kalemleri "Yetersiz geçmiş" olarak ayırır. İnternete bağlanmaz.

Kullanım:
    python main.py                                         # örnek tüketim verisiyle dener
    python main.py --girdi tuketim.xlsx --maliyet maliyetler.xlsx
    python main.py --girdi tuketim.xlsx --abc 70 90 --xyz 0.25 0.5
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import OrderedDict, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from statistics import fmean, pstdev

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
AYLAR = ["ocak", "şubat", "mart", "nisan", "mayıs", "haziran", "temmuz", "ağustos", "eylül", "ekim", "kasım", "aralık"]

POLITIKA = {
    "AX": "Yüksek değer, düzenli talep: sürekli yenileme / tam zamanında tedarik, düşük emniyet stoğu, sık ve küçük sipariş, çerçeve anlaşma.",
    "AY": "Yüksek değer, dalgalı talep: tahmine dayalı planlama, dönemsel gözden geçirme, emniyet stoğunu talep değişkenliğine göre hesaplayın.",
    "AZ": "Yüksek değer, düzensiz talep: mümkünse siparişe göre alım/üretim, stok tutmaktan kaçının, talep sinyallerini yakından izleyin.",
    "BX": "Orta değer, düzenli talep: yeniden sipariş noktası sistemi, ekonomik sipariş miktarı, otomatik sipariş önerisi.",
    "BY": "Orta değer, dalgalı talep: periyodik gözden geçirme, orta düzey emniyet stoğu.",
    "BZ": "Orta değer, düzensiz talep: siparişe göre veya yüksek emniyet stoğuyla; tedarik süresini kısaltmaya çalışın.",
    "CX": "Düşük değer, düzenli talep: basit yeniden sipariş (iki kutu / kanban), toplu alım, seyrek kontrol.",
    "CY": "Düşük değer, dalgalı talep: geniş emniyet stoğu kabul edilebilir, seyrek ve toplu sipariş.",
    "CZ": "Düşük değer, düzensiz talep: stok tutma gerekliliğini sorgulayın; siparişe göre alım veya portföyden çıkarma değerlendirmesi.",
}


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def para(x) -> Decimal:
    if x in (None, ""):
        return Decimal(0)
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("₺", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return Decimal(0)


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def bul(baslik: list, *adlar: str) -> int | None:
    b = [kucuk(x) for x in baslik]
    return next((i for i, x in enumerate(b) if x in adlar), None)


def donem_anahtari(x) -> str | None:
    """Tarih veya dönem metnini 'YYYY-AA' anahtarına çevirir."""
    if isinstance(x, (datetime, date)):
        return f"{x.year:04d}-{x.month:02d}"
    s = kucuk(x)
    m = re.match(r"^(\d{4})[-/.](\d{1,2})", s) or None
    if m:
        return f"{m.group(1)}-{int(m.group(2)):02d}"
    m = re.match(r"^(\d{1,2})[./](\d{1,2})[./](\d{4})$", s)
    if m:
        return f"{m.group(3)}-{int(m.group(2)):02d}"
    m = re.match(r"^(\d{1,2})[./-](\d{4})$", s)
    if m:
        return f"{m.group(2)}-{int(m.group(1)):02d}"
    for i, a in enumerate(AYLAR, 1):
        if s.startswith(a):
            y = re.search(r"(\d{4})", s)
            return f"{y.group(1) if y else '0000'}-{i:02d}"
    return None


def tuketim_oku(yol: Path) -> tuple[dict, dict, dict, list[str]]:
    """Dönüş: (kod → {dönem: miktar}, kod → ad, kod → birim maliyet (dosyada varsa), dönem listesi)."""
    s = tablo_oku(yol)
    b = s[0]
    i_kod = bul(b, "stok kodu", "malzeme kodu", "ürün kodu", "kod", "sku")
    i_ad = bul(b, "stok adı", "malzeme adı", "ürün adı", "açıklama", "ad")
    i_mal = bul(b, "birim maliyet", "birim fiyat", "maliyet")
    i_tar = bul(b, "tarih", "dönem", "ay")
    i_mik = bul(b, "miktar", "tüketim", "çıkış miktarı", "satış miktarı", "adet")
    if i_kod is None:
        raise SystemExit(f"'Stok Kodu' sütunu bulunamadı. Başlıklar: {b}")
    genis = {i: donem_anahtari(x) for i, x in enumerate(b) if i not in (i_kod, i_ad, i_mal) and donem_anahtari(x)}
    veri: dict[str, dict[str, Decimal]] = defaultdict(lambda: defaultdict(Decimal))
    adlar, maliyet = {}, {}
    for r in s[1:]:
        kod = str(r[i_kod] or "").strip()
        if not kod:
            continue
        if i_ad is not None and r[i_ad]:
            adlar[kod] = str(r[i_ad]).strip()
        if i_mal is not None and r[i_mal] not in (None, ""):
            maliyet[kod] = para(r[i_mal])
        if genis and i_mik is None:                    # geniş biçim: her dönem bir sütun
            for i, d in genis.items():
                veri[kod][d] += para(r[i] if i < len(r) else None)
        else:
            if i_tar is None or i_mik is None:
                raise SystemExit("Uzun biçimde Tarih (veya Dönem) ve Miktar sütunları gerekli; geniş biçimde dönem başlıklı sütunlar.")
            d = donem_anahtari(r[i_tar])
            if d:
                veri[kod][d] += para(r[i_mik])
            else:
                veri[kod]                               # kalem listede kalsın
    donemler = sorted({d for v in veri.values() for d in v} | set(genis.values()))
    return veri, adlar, maliyet, donemler


def maliyet_oku(yol: Path | None) -> dict[str, Decimal]:
    if not yol:
        return {}
    s = tablo_oku(yol)
    i_kod = bul(s[0], "stok kodu", "malzeme kodu", "ürün kodu", "kod", "sku")
    i_mal = bul(s[0], "birim maliyet", "birim fiyat", "maliyet")
    if i_kod is None or i_mal is None:
        raise SystemExit(f"Maliyet dosyasında Stok Kodu ve Birim Maliyet gerekli. Başlıklar: {s[0]}")
    return {str(r[i_kod]).strip(): para(r[i_mal]) for r in s[1:] if r[i_kod]}


# ----------------------------------------------------------------------------
# Sınıflandırma
# ----------------------------------------------------------------------------

def abc_sinifla(degerler: dict[str, Decimal], a: float, b: float) -> dict[str, tuple[str, float]]:
    """Kümülatif pay sınırı: bir kalem, kendinden önceki kümülatif pay sınırın altındaysa o sınıfa girer.
    Böylece sınırı aşan ilk kalem de üst sınıfta kalır (ör. %78'den %83'e çıkaran kalem A olur)."""
    toplam = sum(degerler.values(), Decimal(0))
    sonuc, kum = {}, Decimal(0)
    for kod, v in sorted(degerler.items(), key=lambda i: (-i[1], i[0])):
        onceki = float(kum / toplam * 100) if toplam else 100.0
        kum += v
        sinif = "A" if onceki < a else "B" if onceki < b else "C"
        sonuc[kod] = (sinif, float(kum / toplam * 100) if toplam else 0.0)
    return sonuc


def xyz_sinifla(seri: list[float], x: float, y: float) -> tuple[str, float | None]:
    ort = fmean(seri) if seri else 0.0
    if ort <= 0:
        return "-", None
    cv = pstdev(seri) / ort
    return ("X" if cv <= x else "Y" if cv <= y else "Z"), cv


def calistir(girdi: Path, cikti: Path, maliyet_yolu: Path | None = None, abc: tuple[float, float] = (80, 95),
             xyz: tuple[float, float] = (0.5, 1.0), en_az_donem: int = 6) -> dict:
    veri, adlar, maliyet, donemler = tuketim_oku(girdi)
    maliyet.update(maliyet_oku(maliyet_yolu))
    uyarilar = []
    eksik = sorted(k for k in veri if k not in maliyet)
    if eksik:
        uyarilar.append(f"Birim maliyeti olmayan {len(eksik)} kalem değer 0 kabul edildi: {', '.join(eksik[:10])}")
    kalemler = OrderedDict()
    for kod, d in veri.items():
        seri = [float(d.get(p, 0)) for p in donemler]
        toplam = sum(seri)
        # İlk tüketimden önceki dönemler yeni kalem için talep sıfır sayılmaz
        ilk = next((i for i, v in enumerate(seri) if v), None)
        etkin = seri[ilk:] if ilk is not None else []
        kalemler[kod] = {"kod": kod, "ad": adlar.get(kod, ""), "maliyet": maliyet.get(kod, Decimal(0)), "seri": seri,
                         "miktar": toplam, "deger": Decimal(str(toplam)) * maliyet.get(kod, Decimal(0)), "etkin_donem": len(etkin)}
    hareketli = {k: v["deger"] for k, v in kalemler.items() if v["miktar"] > 0}
    abc_sonuc = abc_sinifla(hareketli, *abc)
    toplam_deger = sum(hareketli.values(), Decimal(0))
    for k, v in kalemler.items():
        if v["miktar"] <= 0:
            v.update(abc="-", xyz="-", cv=None, sinif="Hareketsiz", kum=None, pay=0.0,
                     politika="Dönem boyunca tüketim yok: stok yaşlandırma, iade, tasfiye veya değer düşüklüğü değerlendirmesi.")
            continue
        v["abc"], v["kum"] = abc_sonuc[k]
        v["pay"] = float(v["deger"] / toplam_deger * 100) if toplam_deger else 0.0
        seri = v["seri"][len(v["seri"]) - v["etkin_donem"]:]
        v["xyz"], v["cv"] = xyz_sinifla(seri, *xyz)
        if v["etkin_donem"] < en_az_donem:
            v["sinif"] = v["abc"] + "?"
            v["politika"] = f"Yetersiz geçmiş ({v['etkin_donem']} dönem): XYZ güvenilir değil; değer sınıfı {v['abc']}. " \
                            "Yeni kalem ise talebi birkaç dönem izleyin."
        else:
            v["sinif"] = v["abc"] + v["xyz"]
            v["politika"] = POLITIKA[v["sinif"]]
    _rapor(kalemler, donemler, cikti, abc, xyz, uyarilar, toplam_deger)
    return {"kalemler": kalemler, "donemler": donemler, "uyarilar": uyarilar, "toplam_deger": toplam_deger}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
SINIF_DOLGU = {"A": "FDE2E1", "B": "FFF4CE", "C": "E3F5E1"}
PARA = "#,##0.00"


def _rapor(kalemler, donemler, cikti, abc, xyz, uyarilar, toplam_deger):
    wb = Workbook()
    m = wb.active
    m.title = "Matris"
    m.append([f"ABC-XYZ matrisi · ABC: A ≤ %{abc[0]:g}, B ≤ %{abc[1]:g} (kümülatif değer) · XYZ: X ≤ {xyz[0]:g}, Y ≤ {xyz[1]:g} (CV)"
              .replace(".", ",")])
    m["A1"].font = Font(bold=True, size=12)
    m.append([f"{len(donemler)} dönem: {donemler[0] if donemler else '-'} – {donemler[-1] if donemler else '-'} · "
              f"toplam tüketim değeri {toplam_deger:,.2f} TL".replace(",", "X").replace(".", ",").replace("X", ".")])
    m.append([])
    m.append(["Kalem sayısı / değer payı", "X", "Y", "Z", "Toplam"])
    for h in m[4]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for s in "ABC":
        satir = [s]
        for x in "XYZ":
            grup = [v for v in kalemler.values() if v["abc"] == s and v["xyz"] == x and v["sinif"] == s + x]
            satir.append(f"{len(grup)} kalem · %{sum(v['pay'] for v in grup):.1f}".replace(".", ","))
        grup = [v for v in kalemler.values() if v["abc"] == s]
        satir.append(f"{len(grup)} kalem · %{sum(v['pay'] for v in grup):.1f}".replace(".", ","))
        m.append(satir)
        m.cell(m.max_row, 1).fill = PatternFill("solid", fgColor=SINIF_DOLGU[s])
    diger = [("Yetersiz geçmiş", [v for v in kalemler.values() if v["sinif"].endswith("?")]),
             ("Hareketsiz", [v for v in kalemler.values() if v["sinif"] == "Hareketsiz"])]
    for ad, grup in diger:
        m.append([ad, f"{len(grup)} kalem"])
    m.append([])
    m.append(["Sınıf", "Önerilen stok politikası"])
    for h in m[m.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for s, p in POLITIKA.items():
        m.append([s, p])
    for u in uyarilar:
        m.append(["Uyarı", u])
    m.column_dimensions["A"].width = 26
    for c in "BCDE":
        m.column_dimensions[c].width = 22

    k = wb.create_sheet("Kalemler")
    bas = ["Stok Kodu", "Stok Adı", "Birim Maliyet", "Toplam Miktar", "Tüketim Değeri", "Değer Payı %", "Kümülatif %", "ABC",
           "Ort. Dönem Tüketimi", "CV", "XYZ", "Sınıf", "Politika"]
    k.append(bas)
    for h in k[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for v in sorted(kalemler.values(), key=lambda v: (-v["deger"], v["kod"])):
        ort = fmean(v["seri"][len(v["seri"]) - v["etkin_donem"]:]) if v["etkin_donem"] else 0.0
        k.append([v["kod"], v["ad"], float(v["maliyet"]), v["miktar"], float(v["deger"]), v["pay"] / 100,
                  None if v["kum"] is None else v["kum"] / 100, v["abc"], ort, v["cv"], v["xyz"], v["sinif"], v["politika"]])
        r = k.max_row
        for c in (3, 5):
            k.cell(r, c).number_format = PARA
        for c in (6, 7):
            k.cell(r, c).number_format = "0.0%"
        k.cell(r, 9).number_format = "#,##0.0"
        k.cell(r, 10).number_format = "0.00"
        if v["abc"] in SINIF_DOLGU:
            k.cell(r, 8).fill = PatternFill("solid", fgColor=SINIF_DOLGU[v["abc"]])
    for j, w in enumerate((14, 30, 13, 13, 16, 11, 11, 6, 15, 7, 6, 10, 90), 1):
        k.column_dimensions[get_column_letter(j)].width = w
    k.freeze_panes = "C2"
    k.auto_filter.ref = k.dimensions

    # Pareto grafiği (ilk 30 kalem)
    n = min(30, sum(1 for v in kalemler.values() if v["miktar"] > 0))
    if n:
        g = BarChart()
        g.title, g.height, g.width = "Pareto: tüketim değeri ve kümülatif pay", 9, 22
        g.add_data(Reference(k, min_col=5, min_row=1, max_row=1 + n), titles_from_data=True)
        g.set_categories(Reference(k, min_col=1, min_row=2, max_row=1 + n))
        c = LineChart()
        c.add_data(Reference(k, min_col=7, min_row=1, max_row=1 + n), titles_from_data=True)
        c.y_axis.axId = 200
        c.y_axis.crosses = "max"
        g += c
        m.add_chart(g, "G4")

    t = wb.create_sheet("Dönem Tüketimi")
    t.append(["Stok Kodu", "Stok Adı"] + donemler)
    for h in t[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for v in kalemler.values():
        t.append([v["kod"], v["ad"]] + v["seri"])
    t.freeze_panes = "C2"
    b = wb.create_sheet("Bilgi")
    for s in [["ABC", "Kalemler tüketim değerine (miktar × birim maliyet) göre büyükten küçüğe sıralanır; kendinden önceki "
                      "kümülatif pay A sınırının altındaysa A, B sınırının altındaysa B, değilse C"],
              ["XYZ", "CV = dönem tüketimlerinin standart sapması (popülasyon) / ortalaması; kalemin ilk tüketiminden önceki dönemler "
                      "hesaba katılmaz (yeni kalem)"],
              ["Hareketsiz", "Tüm dönemlerde tüketimi sıfır olan kalemler"],
              ["Yetersiz geçmiş", "İlk tüketiminden bu yana belirlenen sayıdan az dönem geçen kalemler (XYZ güvenilir değil)"],
              ["Not", "Sınırlar sektöre göre değişir; --abc ve --xyz ile kendi sınırlarınızı verin"]]:
        b.append(s)
    b.column_dimensions["A"].width = 16
    b.column_dimensions["B"].width = 130
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description="Stokları değer (ABC) ve talep değişkenliği (XYZ) açısından sınıflandırır.")
    ap.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "aylik_tuketim.csv",
                    help="Tüketim (.xlsx/.csv): uzun (Stok Kodu, Tarih/Dönem, Miktar) veya geniş (Stok Kodu, 2026-01, 2026-02 ...)")
    ap.add_argument("--maliyet", type=Path, help="Birim maliyetler (.xlsx/.csv: Stok Kodu, Birim Maliyet); tüketim dosyasında varsa gerekmez")
    ap.add_argument("--abc", nargs=2, type=float, default=[80, 95], metavar=("A", "B"), help="ABC kümülatif pay sınırları, %% (varsayılan 80 95)")
    ap.add_argument("--xyz", nargs=2, type=float, default=[0.5, 1.0], metavar=("X", "Y"), help="XYZ CV sınırları (varsayılan 0.5 1.0)")
    ap.add_argument("--en-az-donem", type=int, default=6, help="XYZ için gereken en az dönem sayısı (varsayılan 6)")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "abc_xyz.xlsx")
    a = ap.parse_args(argv)
    if not (0 < a.abc[0] < a.abc[1] < 100) or not (0 < a.xyz[0] < a.xyz[1]):
        raise SystemExit("Sınırlar artan olmalı: 0 < A < B < 100 ve 0 < X < Y")
    if a.girdi == BURASI / "ornek_veri" / "aylik_tuketim.csv" and not a.maliyet:
        a.maliyet = BURASI / "ornek_veri" / "birim_maliyetler.csv"
    s = calistir(a.girdi, a.cikti, a.maliyet, tuple(a.abc), tuple(a.xyz), a.en_az_donem)
    from collections import Counter
    sayac = Counter(v["sinif"] for v in s["kalemler"].values())
    print(f"[OK] {len(s['kalemler'])} kalem · {len(s['donemler'])} dönem · " + " ".join(f"{k}:{n}" for k, n in sorted(sayac.items())))
    for u in s["uyarilar"]:
        print(f"[!] {u}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
