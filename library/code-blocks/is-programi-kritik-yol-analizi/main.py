"""
İş Programı Kritik Yol Analizi — Workers / Workless kod bloğu
Planlama ve Proje Kontrol › Planlama Mühendisi

Faaliyet süreleri ve bağımlılıklardan kritik yolu, bolluk sürelerini ve proje bitiş tarihini hesaplar.
  İleri geçiş   : erken başlama (ES) = öncüllerin (erken bitiş + gecikme) en büyüğü;   erken bitiş (EF) = ES + süre
  Geri geçiş    : geç bitiş (LF) = ardılların (geç başlama − gecikme) en küçüğü;        geç başlama (LS) = LF − süre
  Toplam bolluk = LS − ES;        serbest bolluk = ardılların (ES − gecikme) en küçüğü − EF  (ardılı yoksa proje bitişi − EF)
  Kritik        = toplam bolluğu 0 olan faaliyet;  proje süresi = en büyük EF
  Tarihler      = --baslangic gününden itibaren; --calisma-gunu 7 (her gün, varsayılan), 6 (pazar hariç) veya 5 (cumartesi-pazar hariç)
Bağımlılıklar yalnız bitiş-başlangıç (FS) türündedir; gecikme "B+2" biçiminde yazılır. Resmi tatiller, kaynak kısıtları ve hava koşulları hesaba katılmaz.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"kod": ("faaliyet kodu", "kod", "faaliyet no", "aktivite kodu", "id"), "ad": ("faaliyet", "faaliyet adi", "is kalemi", "aciklama", "tanim"), "sure": ("sure", "sure gun", "gun"),
         "oncul": ("onculler", "oncul", "onceki faaliyetler", "bagimlilik", "onceki")}


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("kod", "sure")):
        kod = metin(r["kod"])
        if not kod:
            continue
        s = para(r["sure"])
        if s is None or s < 0 or s != int(s):
            raise ValueError(f"{yol.name} satır {r['_satir']}: '{kod}' faaliyetinin süresi okunamadı ('{metin(r['sure'])}'). Süre sıfır veya pozitif tam gün olmalı.")
        if katla(kod) in gorulen:
            raise ValueError(f"{yol.name} satır {r['_satir']}: '{kod}' faaliyet kodu birden fazla satırda.")
        gorulen.add(katla(kod))
        onculler = []
        for p in re.split(r"[,;]", metin(r["oncul"])):
            p = p.strip()
            if not p:
                continue
            m = re.fullmatch(r"(.+?)\s*(?:FS)?\s*([+-]\s*\d+)?", p, re.IGNORECASE)
            onculler.append((m[1].strip(), int((m[2] or "0").replace(" ", ""))))
        l.append({"kod": kod, "ad": metin(r["ad"]), "sure": int(s), "onculler": onculler})
    return l, uy


def gun_tarihi(baslangic: date, n: int, calisma: int = 7) -> date:
    """Başlangıçtan itibaren n. çalışma günü (0 = ilk çalışma günü)."""
    d = baslangic
    while d.weekday() >= calisma:
        d += timedelta(days=1)
    while n > 0:
        d += timedelta(days=1)
        n -= d.weekday() < calisma
    return d


def hesapla(faaliyetler: list[dict], baslangic: date | None = None, calisma: int = 7) -> dict:
    if not faaliyetler:
        raise ValueError("Faaliyet yok.")
    f = {katla(x["kod"]): x for x in faaliyetler}
    for x in faaliyetler:
        x["onc"], x["ard"] = [], []
    for k, x in f.items():
        for kod, lag in x["onculler"]:
            if katla(kod) not in f:
                raise ValueError(f"'{x['kod']}' faaliyetinin öncülü '{kod}' listede yok.")
            if katla(kod) == k:
                raise ValueError(f"'{x['kod']}' faaliyeti kendisinin öncülü olamaz.")
            x["onc"].append((katla(kod), lag))
            f[katla(kod)]["ard"].append((k, lag))
    derece, sira = {k: len(x["onc"]) for k, x in f.items()}, []
    hazir = [k for k, d in derece.items() if d == 0]
    while hazir:
        k = hazir.pop(0)
        sira.append(k)
        for a, _ in f[k]["ard"]:
            derece[a] -= 1
            if derece[a] == 0:
                hazir.append(a)
    if len(sira) < len(f):
        raise ValueError(f"Bağımlılıklarda döngü var: {', '.join(f[k]['kod'] for k in f if k not in sira)}. Bir faaliyet dolaylı olarak kendisini bekliyor.")
    for k in sira:
        x = f[k]
        x["es"] = max((f[p]["ef"] + lag for p, lag in x["onc"]), default=0)
        x["ef"] = x["es"] + x["sure"]
    sure = max(x["ef"] for x in faaliyetler)
    for k in reversed(sira):
        x = f[k]
        x["lf"] = min((f[a]["ls"] - lag for a, lag in x["ard"]), default=sure)
        x["ls"] = x["lf"] - x["sure"]
        x["tb"] = x["ls"] - x["es"]
        x["sb"] = min((f[a]["es"] - lag for a, lag in x["ard"]), default=sure) - x["ef"]
        x["kritik"] = x["tb"] == 0
    for x in faaliyetler:
        x["bas_tarih"] = gun_tarihi(baslangic, x["es"], calisma) if baslangic else None
        x["bit_tarih"] = gun_tarihi(baslangic, max(x["ef"] - 1, x["es"]), calisma) if baslangic else None
    yollar = []

    def izle(k: str, yol: list[str]) -> None:
        devam = [a for a, lag in f[k]["ard"] if f[a]["kritik"] and f[a]["es"] == f[k]["ef"] + lag]
        if not devam and len(yollar) < 10:
            yollar.append(yol)
        for a in devam:
            izle(a, yol + [f[a]["kod"]])

    for k in sira:
        if f[k]["kritik"] and f[k]["es"] == 0 and not any(f[p]["kritik"] and f[p]["ef"] + lag == 0 for p, lag in f[k]["onc"]):
            izle(k, [f[k]["kod"]])
    return {"faaliyetler": sorted(faaliyetler, key=lambda x: (x["es"], x["ef"], x["kod"])), "sure": sure, "yollar": yollar, "baslangic": baslangic, "calisma": calisma,
            "bitis": gun_tarihi(baslangic, max(sure - 1, 0), calisma) if baslangic else None, "kritik_sayisi": sum(x["kritik"] for x in faaliyetler)}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Faaliyetler"
    _baslik(ws, ["Kritik", "Kod", "Faaliyet", "Süre (gün)", "Öncüller", "Erken Başlama", "Erken Bitiş", "Geç Başlama", "Geç Bitiş", "Toplam Bolluk", "Serbest Bolluk", "Başlangıç Tarihi", "Bitiş Tarihi"],
            (8, 10, 34, 9, 18, 10, 10, 10, 10, 10, 10, 13, 13))
    for x in r["faaliyetler"]:
        _satir(ws, ["Evet" if x["kritik"] else "", x["kod"], x["ad"], x["sure"], ", ".join(k + (f"{lag:+d}" if lag else "") for k, lag in x["onculler"]), x["es"], x["ef"], x["ls"], x["lf"], x["tb"], x["sb"],
                    x["bas_tarih"], x["bit_tarih"]], tarih_s=(12, 13), dolgu={1: KIRMIZI if x["kritik"] else None, 10: KIRMIZI if x["kritik"] else None})
    ws.append([])
    takvim = {7: "her gün çalışılıyor", 6: "pazar hariç", 5: "cumartesi ve pazar hariç"}[r["calisma"]]
    ws.append([f"Proje süresi {r['sure']} çalışma günü" + (f"; başlangıç {r['baslangic']:%d.%m.%Y}, bitiş {r['bitis']:%d.%m.%Y} ({takvim})" if r["baslangic"] else "")
               + ". Erken ve geç zamanlar proje başlangıcından itibaren gün sayısıdır (başlangıç = 0). Resmi tatiller ve kaynak kısıtları hesaba katılmaz."])
    ws = wb.create_sheet("Kritik Yol")
    _baslik(ws, ["Yol No", "Kritik Yol", "Faaliyet Sayısı"], (8, 110, 10))
    for i, y in enumerate(r["yollar"], 1):
        _satir(ws, [i, " → ".join(y), len(y)], sar=(2,))
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, baslangic: date | None = None, calisma: int = 7) -> dict:
    faaliyetler, uy = oku(yol)
    r = hesapla(faaliyetler, baslangic, calisma)
    if len(r["yollar"]) > 1:
        uy.append(uyari("Bilgi", "Birden çok kritik yol", yol.name, f"{len(r['yollar'])} kritik yol var; herhangi birindeki gecikme projeyi geciktirir." + (" (ilk 10 yol listelendi)" if len(r["yollar"]) == 10 else "")))
    basi_bos = [x["kod"] for x in r["faaliyetler"] if not x["onc"]]
    sonu_bos = [x["kod"] for x in r["faaliyetler"] if not x["ard"]]
    if len(basi_bos) > 1:
        uy.append(uyari("Bilgi", "Öncülü olmayan birden çok faaliyet", ", ".join(basi_bos), "Hepsi proje başlangıcında başlar. Unutulmuş bir bağımlılık olmadığını kontrol edin."))
    if len(sonu_bos) > 1:
        uy.append(uyari("Bilgi", "Ardılı olmayan birden çok faaliyet", ", ".join(sonu_bos), "Bu faaliyetlerin bolluğu proje bitişine göre hesaplandı. Unutulmuş bir bağımlılık olmadığını kontrol edin."))
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Faaliyet süreleri ve bağımlılıklardan kritik yolu ve bollukları hesaplar.")
    p.add_argument("--faaliyetler", type=Path, default=ORNEK / "faaliyetler.csv", help="Faaliyet Kodu, Faaliyet, Süre (gün), Öncüller (virgülle; gecikme için B+2)")
    p.add_argument("--baslangic", help="Proje başlangıç tarihi GG.AA.YYYY (verilmezse tarih hesaplanmaz; örnek veride 12.10.2026)")
    p.add_argument("--calisma-gunu", choices=("5", "6", "7"), default="7", help="Haftada çalışılan gün: 7 her gün (varsayılan), 6 pazar hariç, 5 hafta sonu hariç")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kritik_yol.xlsx")
    a = p.parse_args(argv)
    bt = a.baslangic or ("12.10.2026" if a.faaliyetler == ORNEK / "faaliyetler.csv" else None)
    bas = tarih(bt) if bt else None
    if bt and bas is None:
        print("[X] --baslangic GG.AA.YYYY biçiminde olmalı")
        return 2
    if not a.faaliyetler.exists():
        print(f"[X] Dosya bulunamadı: {a.faaliyetler}")
        return 1
    try:
        r = calistir(a.faaliyetler, a.cikti, bas, int(a.calisma_gunu))
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(r['faaliyetler'])} faaliyet · proje süresi {r['sure']} gün" + (f" · bitiş {r['bitis']:%d.%m.%Y}" if r["bitis"] else "") + f" · {r['kritik_sayisi']} kritik faaliyet")
    for y in r["yollar"][:3]:
        print(f"[!] Kritik yol: {' → '.join(y)}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
