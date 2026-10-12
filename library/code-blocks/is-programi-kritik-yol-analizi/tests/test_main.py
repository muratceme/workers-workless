import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
PZT = date(2026, 10, 12)


def fa(kod, sure, *onculler):
    return {"kod": kod, "ad": "", "sure": sure, "onculler": [(o, 0) if isinstance(o, str) else o for o in onculler]}


def ag():
    return [fa("A", 5), fa("B", 10, "A"), fa("C", 20, "B"), fa("D", 7, "C"), fa("E", 12, "C"), fa("F", 10, "D", "E"), fa("G", 6, "C"), fa("H", 5, "F", "G")]


def tablo(r):
    return {x["kod"]: (x["es"], x["ef"], x["ls"], x["lf"], x["tb"], x["sb"]) for x in r["faaliyetler"]}


class HesapTesti(unittest.TestCase):
    def test_ileri_ve_geri_gecis(self):
        r = main.hesapla(ag())
        self.assertEqual(tablo(r), {"A": (0, 5, 0, 5, 0, 0), "B": (5, 15, 5, 15, 0, 0), "C": (15, 35, 15, 35, 0, 0), "D": (35, 42, 40, 47, 5, 5), "E": (35, 47, 35, 47, 0, 0), "F": (47, 57, 47, 57, 0, 0),
                                    "G": (35, 41, 51, 57, 16, 16), "H": (57, 62, 57, 62, 0, 0)})
        self.assertEqual((r["sure"], r["yollar"], r["kritik_sayisi"], r["bitis"]), (62, [["A", "B", "C", "E", "F", "H"]], 6, None))

    def test_gecikme_paralel_yollar_ve_bolluk(self):
        r = main.hesapla([fa("A", 3), fa("B", 4, ("A", 2)), fa("C", 2, ("A", -1))])
        self.assertEqual((tablo(r), r["sure"]), ({"A": (0, 3, 0, 3, 0, 0), "C": (2, 4, 7, 9, 5, 5), "B": (5, 9, 5, 9, 0, 0)}, 9))                    # A + 2 gün bekleme; C bir gün bindirmeli
        r = main.hesapla([fa("A", 2), fa("B", 3, "A"), fa("C", 3, "A"), fa("D", 1, "B", "C")])
        self.assertEqual((r["yollar"], r["sure"]), ([["A", "B", "D"], ["A", "C", "D"]], 6))
        r = main.hesapla([fa("A", 2), fa("B", 1, "A"), fa("C", 5, "A"), fa("D", 1, "B"), fa("E", 1, "C", "D")])
        self.assertEqual({k: v[4:] for k, v in tablo(r).items()}, {"A": (0, 0), "B": (3, 0), "C": (0, 0), "D": (3, 3), "E": (0, 0)})                    # B'nin serbest bolluğu 0, toplam bolluğu 3
        r = main.hesapla([fa("M", 0), fa("A", 3, "M"), fa("X", 1)])
        self.assertEqual((tablo(r)["M"], tablo(r)["X"], r["yollar"]), ((0, 0, 0, 0, 0, 0), (0, 1, 2, 3, 2, 2), [["M", "A"]]))                           # sıfır süreli kilometre taşı

    def test_tarihler(self):
        self.assertEqual([main.gun_tarihi(PZT, n, c) for n, c in ((0, 7), (6, 7), (5, 5), (5, 6), (6, 6))], [PZT, date(2026, 10, 18), date(2026, 10, 19), date(2026, 10, 17), date(2026, 10, 19)])
        self.assertEqual((main.gun_tarihi(date(2026, 10, 10), 0, 5), main.gun_tarihi(date(2026, 10, 11), 0, 6)), (PZT, PZT))                           # hafta sonu başlangıcı ilk çalışma gününe kayar
        r = main.hesapla(ag(), PZT)
        self.assertEqual(([(x["kod"], x["bas_tarih"], x["bit_tarih"]) for x in r["faaliyetler"]][:2], r["bitis"]), ([("A", PZT, date(2026, 10, 16)), ("B", date(2026, 10, 17), date(2026, 10, 26))], date(2026, 12, 12)))
        r = main.hesapla(ag(), PZT, 5)
        self.assertEqual(([(x["bas_tarih"], x["bit_tarih"]) for x in r["faaliyetler"]][1], r["bitis"]), ((date(2026, 10, 19), date(2026, 10, 30)), date(2027, 1, 5)))        # 62 iş günü

    def test_hatalar(self):
        for kotu in ([], [fa("A", 1, "Z")], [fa("A", 1, "A")], [fa("A", 1, "C"), fa("B", 1, "A"), fa("C", 1, "B")]):
            with self.assertRaises(ValueError):
                main.hesapla(kotu)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "k.xlsx"
            r = main.calistir(ORNEK / "faaliyetler.csv", cikti, PZT)
            self.assertEqual((r["sure"], r["bitis"], r["yollar"], r["uyarilar"]), (62, date(2026, 12, 12), [["A", "B", "C", "E", "F", "H"]], []))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Faaliyetler", "Kritik Yol", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "f.csv"
            y.write_text("Kod;Faaliyet Adı;Süre;Öncül\nA;Kazı;5;\nb;Temel;10;a\nC;Kalıp;3;A + 2; \nD;Beton;1;B, C FS-1\nE;Kür;7;D;\n", encoding="utf-8")
            l, _ = main.oku(y)
            self.assertEqual([(x["kod"], x["sure"], x["onculler"]) for x in l], [("A", 5, []), ("b", 10, [("a", 0)]), ("C", 3, [("A", 2)]), ("D", 1, [("B", 0), ("C", -1)]), ("E", 7, [("D", 0)])])
            c = str(Path(tmp) / "x.xlsx")
            r = main.calistir(y, Path(c))
            self.assertEqual((r["sure"], r["yollar"], r["bitis"]), (23, [["A", "b", "D", "E"]], None))
            arg = ["--faaliyetler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--baslangic", "12.10.2026", "--calisma-gunu", "6"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(arg + ["--baslangic", "yarın"]), 2)
            self.assertEqual(main.main(["--faaliyetler", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Kod;Süre;Öncül\nA;beş;\n", "Kod;Süre;Öncül\nA;1;\nA;2;\n", "Kod;Süre;Öncül\nA;1;B\nB;1;A\n", "Kod;Süre\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            y.write_text("Kod;Süre;Öncül\nA;1;\nB;1;\nC;1;A\n", encoding="utf-8")
            self.assertEqual([u["tur"] for u in main.calistir(y, Path(c))["uyarilar"]], ["Ardılı olmayan birden çok faaliyet", "Öncülü olmayan birden çok faaliyet"])
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
