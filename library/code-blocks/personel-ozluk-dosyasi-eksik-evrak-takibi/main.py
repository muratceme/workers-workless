"""
Personel Özlük Dosyası Eksik Evrak Takibi — Workers / Workless kod bloğu
İnsan Kaynakları › İnsan Kaynakları Uzman Yardımcısı

Çalışan listesini zorunlu özlük evrakı listesiyle karşılaştırır, eksik ve süresi dolan evrakı çalışan bazında
raporlar:
  - Evrak listesi koşullu olabilir: Tümü, Cinsiyet=Erkek, Uyruk!=TC, Pozisyon~şoför, Tehlike Sınıfı=Çok Tehlikeli,
    Gece Çalışması=Evet… Aynı evrak farklı koşullarla farklı geçerlilik süresine sahip olabilir (ilk uyan satır).
  - Teslim süresi: işe girişten itibaren evrakın verilmesi gereken gün sayısı (0 = işe başlamadan önce; boş = 30).
    Süre dolmadıysa "Bekleniyor", dolduysa "Eksik".
  - Geçerlilik (ay): teslim tarihinden itibaren; geçmişse "Süresi dolmuş", 30 gün içinde dolacaksa "Yaklaşıyor".
  - Listede olmayan evrak adı ve çalışan listesinde olmayan sicil kayıtları ayrıca bildirilir.
Rapor: evrak matrisi (çalışan × evrak), eksik listesi, departman özeti, çalışana hatırlatma metinleri. İnternete
bağlanmaz.

Kullanım:
    python main.py                                                   # örnek: 8 çalışan, 18 evrak türü, durum tarihi 08.10.2026
    python main.py --calisanlar calisanlar.xlsx --kayitlar evrak_kayitlari.xlsx --liste evrak_listesi.csv --bugun 08.10.2026
"""
from __future__ import annotations

import argparse
import calendar
import csv
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
VARSAYILAN_TESLIM = 30

LISTE_SUTUNLARI = {"evrak": ("evrak", "belge", "evrak adi"), "kosul": ("kosul", "kimler", "kapsam"), "gecerlilik": ("gecerlilik ay", "gecerlilik"),
                   "teslim": ("teslim suresi gun", "teslim suresi", "teslim gun")}
CALISAN_SUTUNLARI = {"sicil": ("sicil no", "sicil", "personel no"), "ad": ("ad soyad", "ad"), "departman": ("departman", "birim"),
                     "pozisyon": ("pozisyon", "unvan", "gorev"), "giris": ("ise giris", "ise giris tarihi", "giris tarihi")}
KAYIT_SUTUNLARI = {"sicil": ("sicil no", "sicil", "personel no"), "evrak": ("evrak", "belge", "evrak adi"), "teslim": ("teslim tarihi", "tarih", "belge tarihi"),
                   "bitis": ("gecerlilik bitis", "gecerlilik bitis tarihi", "son gecerlilik")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def ay_ekle(d: date, ay: int) -> date:
    y, m = divmod(d.month - 1 + ay, 12)
    yil, ay_ = d.year + y, m + 1
    return date(yil, ay_, min(d.day, calendar.monthrange(yil, ay_)[1]))


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> tuple[list[dict], list[str]]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            basliklar = [metin(c) for c in r]
            return ([{a: (r2[j] if j < len(r2) else None) for a, j in es.items()} | {"_satir": n, "_ham": dict(zip(b, r2))}
                     for n, r2 in enumerate(s[bi + 1:], bi + 2)], basliklar)
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Gereklilik:
    evrak: str
    kosul: str
    gecerlilik: int | None
    teslim: int

    def uyar(self, c: "Calisan") -> bool:
        k = self.kosul.strip()
        if not k or katla(k) in ("tumu", "hepsi", "herkes"):
            return True
        m = re.match(r"\s*(.+?)\s*(!=|=|~)\s*(.+)\s*$", k)
        if not m:
            return False
        alan, op, deger = katla(m.group(1)), m.group(2), katla(m.group(3))
        v = katla(c.alanlar.get(alan, ""))
        return v == deger if op == "=" else (v != deger and v != "") if op == "!=" else deger in v


@dataclass
class Calisan:
    sicil: str
    ad: str
    departman: str
    pozisyon: str
    giris: date | None
    alanlar: dict
    durumlar: dict = field(default_factory=dict)   # evrak → (durum, açıklama, tarih)


def oku_liste(yol: Path) -> list[Gereklilik]:
    rows, _ = kayitlar(yol, LISTE_SUTUNLARI, ("evrak",))
    sonuc = []
    for r in rows:
        if r.get("evrak"):
            g, t = metin(r.get("gecerlilik")), metin(r.get("teslim"))
            sonuc.append(Gereklilik(metin(r["evrak"]), metin(r.get("kosul")), int(float(g.replace(",", "."))) if g else None,
                                    int(float(t.replace(",", "."))) if t else VARSAYILAN_TESLIM))
    return sonuc


def oku_calisanlar(yol: Path) -> list[Calisan]:
    rows, _ = kayitlar(yol, CALISAN_SUTUNLARI, ("sicil",))
    return [Calisan(metin(r["sicil"]), metin(r.get("ad")), metin(r.get("departman")) or "—", metin(r.get("pozisyon")), tarih(r.get("giris")),
                    {k: metin(v) for k, v in r["_ham"].items() if k} | {"pozisyon": metin(r.get("pozisyon"))})
            for r in rows if r.get("sicil")]


def oku_kayitlar(yol: Path) -> list[dict]:
    rows, _ = kayitlar(yol, KAYIT_SUTUNLARI, ("sicil", "evrak"))
    return [{"sicil": metin(r["sicil"]), "evrak": metin(r["evrak"]), "teslim": tarih(r.get("teslim")), "bitis": tarih(r.get("bitis")), "satir": r["_satir"]}
            for r in rows if r.get("sicil") and r.get("evrak")]


def evrak_eslestir(ad: str, adlar: list[str]) -> str | None:
    k = katla(ad)
    for a in adlar:
        if katla(a) == k:
            return a
    for a in adlar:
        ka = katla(a)
        if ka and (ka in k or k in ka):
            return a
    return None


# ----------------------------------------------------------------------------
# Analiz
# ----------------------------------------------------------------------------

def analiz_et(calisanlar: list[Calisan], liste: list[Gereklilik], kayit_listesi: list[dict], bugun: date, uyari_gun: int = 30) -> dict:
    uyarilar = []
    evraklar = list(dict.fromkeys(g.evrak for g in liste))
    sicil = {c.sicil: c for c in calisanlar}
    teslimler = defaultdict(list)
    for k in kayit_listesi:
        if k["sicil"] not in sicil:
            uyarilar.append({"onem": "Orta", "sicil": k["sicil"], "aciklama": f"'{k['evrak']}' kaydı var ama {k['sicil']} çalışan listesinde yok (ayrılmış olabilir)"})
            continue
        ad = evrak_eslestir(k["evrak"], evraklar)
        if ad is None:
            uyarilar.append({"onem": "Bilgi", "sicil": k["sicil"], "aciklama": f"'{k['evrak']}' evrak listesinde yok; adını listeyle eşleştirin"})
            continue
        teslimler[(k["sicil"], ad)].append(k)
    eksikler = []
    for c in calisanlar:
        for ad in evraklar:
            g = next((x for x in liste if x.evrak == ad and x.uyar(c)), None)
            if g is None:
                c.durumlar[ad] = ("Gerekmez", "", None)
                continue
            kayit = max(teslimler.get((c.sicil, ad), []), key=lambda k: k["teslim"] or date.min, default=None)
            if kayit is None:
                son = c.giris + timedelta(days=g.teslim) if c.giris else None
                if son and bugun <= son:
                    c.durumlar[ad] = ("Bekleniyor", f"son teslim {son:%d.%m.%Y}", son)
                else:
                    c.durumlar[ad] = ("Eksik", f"son teslim {son:%d.%m.%Y}" if son else "işe giriş tarihi yok", son)
                    eksikler.append({"c": c, "evrak": ad, "durum": "Eksik", "onem": "Yüksek" if g.teslim == 0 else "Orta",
                                     "aciklama": "İşe başlamadan önce alınmalıydı" if g.teslim == 0 else f"İşe girişten itibaren {g.teslim} gün içinde alınmalıydı"})
                continue
            bitis = kayit["bitis"] or (ay_ekle(kayit["teslim"], g.gecerlilik) if g.gecerlilik and kayit["teslim"] else None)
            if bitis and bitis < bugun:
                c.durumlar[ad] = ("Süresi dolmuş", f"{bitis:%d.%m.%Y} doldu", bitis)
                eksikler.append({"c": c, "evrak": ad, "durum": "Süresi dolmuş", "onem": "Yüksek",
                                 "aciklama": f"{kayit['teslim']:%d.%m.%Y} tarihli belge, {g.gecerlilik} ay geçerli; {bitis:%d.%m.%Y} tarihinde doldu"
                                 if g.gecerlilik and not kayit["bitis"] else f"Geçerlilik {bitis:%d.%m.%Y} tarihinde doldu"})
            elif bitis and (bitis - bugun).days <= uyari_gun:
                c.durumlar[ad] = ("Yaklaşıyor", f"{bitis:%d.%m.%Y} dolacak", bitis)
                eksikler.append({"c": c, "evrak": ad, "durum": "Yaklaşıyor", "onem": "Orta",
                                 "aciklama": f"Geçerlilik {bitis:%d.%m.%Y} tarihinde dolacak ({(bitis - bugun).days} gün)"})
            else:
                c.durumlar[ad] = ("Tamam", f"{kayit['teslim']:%d.%m.%Y}" if kayit["teslim"] else "tarih yok", bitis)
    ozet = []
    dept = defaultdict(list)
    for c in calisanlar:
        dept[c.departman].append(c)
    for d, lst in sorted(dept.items()):
        gerekli = [s for c in lst for s, *_ in c.durumlar.values() if s != "Gerekmez"]
        ozet.append({"departman": d, "calisan": len(lst), "gerekli": len(gerekli), "tamam": gerekli.count("Tamam"), "eksik": gerekli.count("Eksik"),
                     "dolmus": gerekli.count("Süresi dolmuş"), "yaklasan": gerekli.count("Yaklaşıyor"), "bekleniyor": gerekli.count("Bekleniyor"),
                     "tam_dosya": sum(1 for c in lst if all(s in ("Tamam", "Gerekmez", "Bekleniyor") for s, *_ in c.durumlar.values()))})
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    eksikler.sort(key=lambda e: (sira[e["onem"]], e["c"].sicil))
    return {"evraklar": evraklar, "eksikler": eksikler, "ozet": ozet, "uyarilar": uyarilar}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Tamam": "E3F4E1", "Eksik": "FDE2E1", "Süresi dolmuş": "FDE2E1", "Yaklaşıyor": "FFF4CE", "Bekleniyor": "E8F0FE", "Gerekmez": "F2F2F2",
        "Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KISA = {"Tamam": "✓", "Eksik": "EKSİK", "Süresi dolmuş": "SÜRE DOLDU", "Yaklaşıyor": "Yaklaşıyor", "Bekleniyor": "Bekleniyor", "Gerekmez": "—"}
KONTROL = PatternFill("solid", fgColor="FFF4CE")
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def hatirlatma(c: Calisan, satirlar: list[dict]) -> str:
    maddeler = "\n".join(f"- {e['evrak']}" + (" (süresi dolmuş, yenisi gerekli)" if e["durum"] == "Süresi dolmuş" else
                                             " (geçerliliği yakında doluyor)" if e["durum"] == "Yaklaşıyor" else "") for e in satirlar)
    return (f"Sayın {c.ad},\n\nÖzlük dosyanızın tamamlanması için aşağıdaki belgelerin İnsan Kaynakları birimine iletilmesini rica ederiz:\n\n"
            f"{maddeler}\n\nBelgeleri en geç [TARİH] tarihine kadar teslim edebilirsiniz. Sorularınız için bize ulaşabilirsiniz.\n\nİnsan Kaynakları")


def rapor_yaz(cikti: Path, calisanlar: list[Calisan], s: dict, bugun: date) -> None:
    wb = Workbook()
    mx = wb.active
    mx.title = "Evrak Matrisi"
    _baslik(mx, ["Sicil", "Ad Soyad", "Departman", "Pozisyon", "İşe Giriş"] + s["evraklar"] + ["Eksik / Dolmuş"], (8, 18, 14, 18, 11) + (11,) * len(s["evraklar"]) + (10,))
    mx.row_dimensions[1].height = 60
    for c in calisanlar:
        mx.append([c.sicil, c.ad, c.departman, c.pozisyon, c.giris] + [KISA[c.durumlar[e][0]] for e in s["evraklar"]]
                  + [sum(1 for d, *_ in c.durumlar.values() if d in ("Eksik", "Süresi dolmuş"))])
        mx.cell(mx.max_row, 5).number_format = "DD.MM.YYYY"
        for j, e in enumerate(s["evraklar"], 6):
            mx.cell(mx.max_row, j).fill = PatternFill("solid", fgColor=RENK[c.durumlar[e][0]])
            mx.cell(mx.max_row, j).alignment = Alignment(horizontal="center")
    mx.append([])
    mx.append([f"Durum tarihi {bugun:%d.%m.%Y} · ✓ tamam · EKSİK · SÜRE DOLDU · Yaklaşıyor (30 gün içinde) · Bekleniyor (teslim süresi dolmadı) · — gerekmez"])

    ek = wb.create_sheet("Eksik Listesi")
    _baslik(ek, ["Önem", "Sicil", "Ad Soyad", "Departman", "Evrak", "Durum", "Açıklama", "Talep Edildi", "Teslim Alındı"], (9, 8, 18, 14, 30, 13, 60, 12, 12))
    for e in s["eksikler"]:
        c = e["c"]
        ek.append([e["onem"], c.sicil, c.ad, c.departman, e["evrak"], e["durum"], e["aciklama"], "", ""])
        ek.cell(ek.max_row, 1).fill = PatternFill("solid", fgColor=RENK[e["onem"]])
        ek.cell(ek.max_row, 6).fill = PatternFill("solid", fgColor=RENK[e["durum"]])
        ek.cell(ek.max_row, 8).fill = ek.cell(ek.max_row, 9).fill = KONTROL
    ek.auto_filter.ref = f"A1:I{ek.max_row}"

    dz = wb.create_sheet("Departman Özeti")
    _baslik(dz, ["Departman", "Çalışan", "Gerekli Evrak", "Tamam", "Eksik", "Süresi Dolmuş", "Yaklaşan", "Bekleniyor", "Tamamlanma", "Tam Dosyalı Çalışan"],
            (16, 8, 10, 8, 8, 10, 9, 10, 11, 12))
    for x in s["ozet"]:
        dz.append([x["departman"], x["calisan"], x["gerekli"], x["tamam"], x["eksik"], x["dolmus"], x["yaklasan"], x["bekleniyor"],
                   x["tamam"] / x["gerekli"] if x["gerekli"] else None, x["tam_dosya"]])
        dz.cell(dz.max_row, 9).number_format = "0%"

    ht = wb.create_sheet("Hatırlatma Metinleri")
    _baslik(ht, ["Sicil", "Ad Soyad", "Metin (taslak)"], (8, 18, 100))
    gruplu = defaultdict(list)
    for e in s["eksikler"]:
        gruplu[e["c"].sicil].append(e)
    for c in calisanlar:
        if gruplu.get(c.sicil):
            ht.append([c.sicil, c.ad, hatirlatma(c, gruplu[c.sicil])])
            ht.cell(ht.max_row, 3).alignment = UST

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Sicil", "Açıklama"], (9, 8, 90))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["sicil"], u["aciklama"]])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(calisan_yolu: Path, kayit_yolu: Path, liste_yolu: Path, cikti: Path, bugun: date | None = None, uyari_gun: int = 30) -> dict:
    bugun = bugun or date.today()
    calisanlar = oku_calisanlar(calisan_yolu)
    s = analiz_et(calisanlar, oku_liste(liste_yolu), oku_kayitlar(kayit_yolu), bugun, uyari_gun)
    rapor_yaz(cikti, calisanlar, s, bugun)
    return {**s, "calisanlar": calisanlar, "bugun": bugun}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Çalışan özlük dosyalarını zorunlu evrak listesiyle karşılaştırıp eksik ve süresi dolan belgeleri raporlar.")
    p.add_argument("--calisanlar", type=Path, default=ORNEK / "calisanlar.csv", help="Sicil No, Ad Soyad, Departman, Pozisyon, İşe Giriş + koşul sütunları")
    p.add_argument("--kayitlar", type=Path, default=ORNEK / "evrak_kayitlari.csv", help="Sicil No, Evrak, Teslim Tarihi, [Geçerlilik Bitiş]")
    p.add_argument("--liste", type=Path, default=ORNEK / "evrak_listesi.csv", help="Evrak, Koşul, Geçerlilik (ay), Teslim Süresi (gün)")
    p.add_argument("--bugun", help="Durum tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 08.10.2026)")
    p.add_argument("--uyari-gun", type=int, default=30, help="Geçerliliği bu kadar gün içinde dolacak belgeler uyarılır (varsayılan 30)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "ozluk_eksik_evrak.xlsx")
    a = p.parse_args(argv)
    for y in (a.calisanlar, a.kayitlar, a.liste):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    bugun = tarih(a.bugun) if a.bugun else (date(2026, 10, 8) if a.calisanlar == ORNEK / "calisanlar.csv" else None)
    if a.bugun and bugun is None:
        print("[X] --bugun GG.AA.YYYY biçiminde olmalı")
        return 1
    try:
        s = calistir(a.calisanlar, a.kayitlar, a.liste, a.cikti, bugun, a.uyari_gun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(s['calisanlar'])} çalışan · {len(s['evraklar'])} evrak türü · {len(s['eksikler'])} eksik / süreli kayıt · {s['bugun']:%d.%m.%Y}")
    for e in s["eksikler"]:
        if e["onem"] == "Yüksek":
            print(f"[!] {e['c'].sicil} {e['c'].ad}: {e['evrak']} — {e['durum']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
