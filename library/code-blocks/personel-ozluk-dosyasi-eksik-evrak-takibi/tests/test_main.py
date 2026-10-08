import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 8)


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "o.xlsx"
        cls.s = main.calistir(ORNEK / "calisanlar.csv", ORNEK / "evrak_kayitlari.csv", ORNEK / "evrak_listesi.csv", cls.cikti, BUGUN)
        cls.c = {c.sicil: c for c in cls.s["calisanlar"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def durum(self, sicil, evrak):
        return self.c[sicil].durumlar[evrak][0]

    def test_kosullar(self):
        self.assertEqual(self.durum("1003", "Askerlik durum belgesi"), "Gerekmez")      # kadın
        self.assertEqual(self.durum("1006", "Çalışma izni"), "Eksik")                    # yabancı uyruk
        self.assertEqual(self.durum("1001", "Çalışma izni"), "Gerekmez")
        self.assertEqual(self.durum("1004", "Sürücü belgesi ve SRC belgesi"), "Tamam")  # şoför
        self.assertEqual(self.durum("1008", "Sürücü belgesi ve SRC belgesi"), "Gerekmez")
        self.assertEqual(self.durum("1003", "Gece çalışması sağlık raporu"), "Gerekmez")

    def test_gecerlilik(self):
        self.assertEqual(self.durum("1002", "Periyodik sağlık muayenesi"), "Süresi dolmuş")   # tehlikeli 36 ay
        self.assertEqual(self.durum("1003", "Periyodik sağlık muayenesi"), "Tamam")           # az tehlikeli 60 ay
        self.assertEqual(self.durum("1002", "Temel İSG eğitimi belgesi"), "Yaklaşıyor")
        self.assertEqual(self.c["1007"].durumlar["Temel İSG eğitimi belgesi"][2], date(2026, 4, 10))

    def test_teslim_suresi(self):
        bekleyen = [e for e, (d, *_) in self.c["1005"].durumlar.items() if d == "Bekleniyor"]
        self.assertEqual(len(bekleyen), 6)
        self.assertEqual(self.durum("1005", "Gece çalışması sağlık raporu"), "Eksik")
        e = next(x for x in self.s["eksikler"] if x["c"].sicil == "1005")
        self.assertEqual((e["onem"], e["aciklama"]), ("Yüksek", "İşe başlamadan önce alınmalıydı"))
        self.assertEqual(sum(1 for x in self.s["eksikler"] if x["onem"] == "Yüksek"), 9)

    def test_uyarilar_ve_ozet(self):
        u = [x["aciklama"] for x in self.s["uyarilar"]]
        self.assertTrue(any("Ehliyet fotokopisi" in x for x in u))
        self.assertTrue(any("1009" in x for x in u))
        oz = {x["departman"]: x for x in self.s["ozet"]}
        self.assertEqual((oz["Lojistik"]["tam_dosya"], oz["Üretim"]["bekleniyor"]), (1, 6))

    def test_excel(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Evrak Matrisi", "Eksik Listesi", "Departman Özeti", "Hatırlatma Metinleri", "Uyarılar"])
        mx = wb["Evrak Matrisi"]
        self.assertEqual(mx.max_column, 5 + 16 + 1)
        self.assertIn("Çalışma izni", wb["Hatırlatma Metinleri"]["C6"].value)


class KuralTesti(unittest.TestCase):
    def test_yardimcilar(self):
        self.assertEqual(main.ay_ekle(date(2024, 1, 31), 1), date(2024, 2, 29))
        self.assertEqual(main.evrak_eslestir("kimlik fotokopisi (ön-arka)", ["Kimlik fotokopisi", "Diploma"]), "Kimlik fotokopisi")
        c = main.Calisan("1", "A", "D", "Şoför", None, {"uyruk": "", "pozisyon": "Tır Şoförü"})
        self.assertFalse(main.Gereklilik("X", "Uyruk!=TC", None, 0).uyar(c))
        self.assertTrue(main.Gereklilik("X", "Pozisyon~şoför", None, 0).uyar(c))

    def test_bitis_tarihi_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t / "c.csv").write_text("Sicil No;Ad Soyad;İşe Giriş\n1;A;01.01.2026\n", encoding="utf-8")
            (t / "l.csv").write_text("Evrak;Koşul;Geçerlilik (ay)\nPasaport;Tümü;\n", encoding="utf-8")
            (t / "k.csv").write_text("Sicil No;Evrak;Teslim Tarihi;Geçerlilik Bitiş\n1;Pasaport;01.01.2026;01.10.2026\n", encoding="utf-8")
            s = main.calistir(t / "c.csv", t / "k.csv", t / "l.csv", t / "o.xlsx", BUGUN)
            self.assertEqual(s["calisanlar"][0].durumlar["Pasaport"][0], "Süresi dolmuş")
            self.assertEqual(main.main(["--cikti", str(t / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--liste", str(t / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
