import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def u(fiyat, eski, yeni, hedef=None, asgari=None, adet=None, kod="K", kategori="G"):
    d = lambda x: None if x is None else D(str(x))   # noqa: E731
    return {"kod": kod, "ad": "Ad", "kategori": kategori, "fiyat": d(fiyat), "eski": d(eski), "yeni": d(yeni), "hedef": d(hedef), "asgari": d(asgari), "adet": d(adet)}


def g(urun, **kw):
    r = main.guncelle([urun], **kw)
    x = r["urunler"][0]
    return x["yeni_fiyat"], x["degisim"], x["yeni_marj"], x["durum"], x["nota"], [(w["onem"], w["tur"]) for w in r["uyarilar"]]


class HesapTesti(unittest.TestCase):
    def test_hedef_ve_yuvarlama(self):
        self.assertEqual(g(u(100, 70, 77)), (D("110.00"), D("10.00"), D("30.00"), "Artış", "", []))
        self.assertEqual([g(u(100, 70, 77, 35), basamak=D(b))[0] for b in ("0.01", "0.5", "5")], [D("118.47"), D("118.5"), D(120)])
        self.assertEqual((g(u(100, 70, 77), marj=D(50))[:3], g(u(100, 70, 77, 35), marj=D(50))[0], g(u(100, None, 77), marj=D(50))[0]), ((D("154.00"), D("54.00"), D("50.00")), D("118.47"), D("154.00")))
        self.assertEqual((g(u(100, 70, 77, 40), yontem="katki")[0], g(u(100, 70, 77), yontem="katki")[0], g(u(100, 80, 100, 100), yontem="katki")[0]), (D("107.80"), D("110.00"), D("200.00")))
        self.assertEqual(g(u(45, 30, 31), basamak=D("0.5"))[:2], (D("46.5"), D("3.33")))
        x = main.guncelle([u(100, 70, 77, 35)])["urunler"][0]
        self.assertEqual((x["eski_marj"], x["degismezse"], x["hedef_kullanilan"], x["fark"], x["ciro_etkisi"]), (D("30.00"), D("23.00"), D("35.00"), D("18.47"), None))
        self.assertEqual(g(u(100, 70, None)), (D(100), D("0.00"), D("30.00"), "Aynı", "Yeni maliyet yok", [("Orta", "Yeni maliyet yok")]))
        self.assertEqual([g(x)[3:] for x in (u(100, None, 77), u(100, 70, 77, 100))], [("Aynı", "Hedef belirlenemedi", [("Orta", "Hedef marj belirlenemedi")])] * 2)
        self.assertEqual(g(u(100, 70, 110, -10)), (D("100.00"), D("0.00"), D("-10.00"), "Aynı", "", [("Yüksek", "Maliyetin altında fiyat")]))

    def test_kurallar(self):
        self.assertEqual(g(u(100, 70, 63)), (D(100), D("0.00"), D("37.00"), "Aynı", "İndirim yapılmadı", [("Bilgi", "Maliyet düştü, fiyat korundu")]))
        self.assertEqual(g(u(100, 70, 63), indirim=True), (D("90.00"), D("-10.00"), D("30.00"), "İndirim", "", []))
        self.assertEqual(g(u(100, 70, 84), azami=D(20)), (D("120.00"), D("20.00"), D("30.00"), "Artış", "", []))
        self.assertEqual(g(u(100, 70, "84.07"), azami=D(20)), (D("120.00"), D("20.00"), D("29.94"), "Artış", "Artış tavanı (%20) uygulandı", [("Orta", "Hedef marjın altında")]))
        self.assertEqual(g(u(33, 20, 30), azami=D(10), basamak=D("0.5"))[0], D("36.0"))
        self.assertEqual(g(u(100, 70, 70, asgari=105)), (D(105), D("5.00"), D("33.33"), "Artış", "Asgari fiyat uygulandı", [("Bilgi", "Asgari fiyat uygulandı")]))
        self.assertEqual(g(u(100, 70, 140, asgari=130), azami=D(10))[:3:2], (D(130), D("-7.69")))
        self.assertEqual(g(u(100, 70, 70, asgari=100)), (D("100.00"), D("0.00"), D("30.00"), "Aynı", "", []))
        self.assertEqual(g(u(100, 70, "71.4"), asgari_degisim=D(3)), (D(100), D("0.00"), D("28.60"), "Aynı", "Değişim %3 eşiğinin altında", [("Orta", "Hedef marjın altında")]))
        self.assertEqual(g(u(100, 70, "71.4"), asgari_degisim=D(2))[:2], (D("102.00"), D("2.00")))
        self.assertEqual(g(u(40, 35, 52), azami=D(20), basamak=D("0.5")), (D("48.0"), D("20.00"), D("-8.33"), "Artış", "Artış tavanı (%20) uygulandı", [("Yüksek", "Maliyetin altında fiyat")]))
        with self.assertRaises(ValueError):
            main.guncelle([])

    def test_ozet(self):
        r = main.guncelle([u(100, 70, 77, adet=10, kod="1", kategori="A"), u(100, 70, 63, adet=5, kod="2", kategori="a"), u(50, 25, 30, kod="3", kategori="B"), u(200, 100, 90, adet=2, kod="4", kategori="B")], indirim=True)
        self.assertEqual(r["ozet"], {"urun": 4, "artis": 2, "indirim": 2, "ayni": 0, "ort_degisim": D("2.50"), "ciro_etkisi": D("10.00"), "ciro_etkisi_oran": D("0.53")})
        self.assertEqual(r["kategoriler"], [{"kategori": "A", "urun": 2, "artis": 1, "indirim": 1, "ayni": 0, "ort_degisim": D("0.00"), "ciro_etkisi": D("50.00"), "ciro_etkisi_oran": D("3.33")},
                                            {"kategori": "B", "urun": 2, "artis": 1, "indirim": 1, "ayni": 0, "ort_degisim": D("5.00"), "ciro_etkisi": D("-40.00"), "ciro_etkisi_oran": D("-10.00")}])
        self.assertEqual(main.guncelle([u(100, 70, 77)])["ozet"]["ciro_etkisi"], None)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "urunler.csv", cikti, basamak=D("0.5"), azami=D(20))
            self.assertEqual([(x["kod"], x["eski_marj"], x["degismezse"], x["yeni_fiyat"], x["degisim"], x["yeni_marj"], x["durum"], x["ciro_etkisi"]) for x in r["urunler"]],
                             [("U-01", D("30.00"), D("23.00"), D(110), D("10.00"), D("30.00"), "Artış", D(5000)), ("U-02", D("25.00"), D("25.00"), D(60), D("0.00"), D("25.00"), "Aynı", D(0)),
                              ("U-03", D("40.00"), D("22.00"), D(300), D("20.00"), D("35.00"), "Artış", D(4000)), ("U-04", D("33.33"), D("40.00"), D(180), D("0.00"), D("40.00"), "Aynı", D(0)),
                              ("U-05", D("33.33"), D("31.11"), D("46.5"), D("3.33"), D("33.33"), "Artış", D(600)), ("U-06", D("25.00"), D("10.00"), D(500), D("25.00"), D("28.00"), "Artış", D(6000)),
                              ("U-07", D("6.67"), D("-10.00"), D("35.5"), D("18.33"), D("7.04"), "Artış", D(1100)), ("U-08", D("33.33"), None, D(90), D("0.00"), D("33.33"), "Aynı", None),
                              ("U-09", D("12.50"), D("-30.00"), D(48), D("20.00"), D("-8.33"), "Artış", D(1200))])
            self.assertEqual(r["ozet"], {"urun": 9, "artis": 6, "indirim": 0, "ayni": 3, "ort_degisim": D("10.74"), "ciro_etkisi": D(17900), "ciro_etkisi_oran": D("10.94")})
            self.assertEqual([(k["kategori"], k["urun"], k["artis"], k["ort_degisim"], k["ciro_etkisi"]) for k in r["kategoriler"]],
                             [("Bağlantı", 2, 1, D("5.00"), D(5000)), ("Boya", 3, 3, D("15.56"), D(7700)), ("El Aleti", 4, 2, D("10.00"), D(5200))])
            self.assertEqual([(w["onem"], w["tur"], w["kim"]) for w in r["uyarilar"]],
                             [("Yüksek", "Maliyetin altında fiyat", "U-09 · Maket Bıçağı"), ("Orta", "Hedef marjın altında", "U-03 · Matkap Ucu Seti"), ("Orta", "Yeni maliyet yok", "U-08 · Çelik Metre 5 m"),
                              ("Bilgi", "Asgari fiyat uygulandı", "U-06 · Astar Boya 2,5 L"), ("Bilgi", "Maliyet düştü, fiyat korundu", "U-04 · Silikon Tabancası")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Yeni Fiyat Listesi", "Değişim Raporu", "Özet", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "u.csv"
            y.write_text("Stok Kodu;Stok Adı;Grup;Liste Fiyatı;Önceki Maliyet;Güncel Maliyet;Hedef Oran;Taban Fiyat;Satış Adedi\nA1;Ürün;;1.250,50;900;950,25;32,5;1000;12\na1;Kopya;;5;;;;;\nA2;X;G;0;1;1;;;\nA3;X;G;5;x;1;;;\n"
                         "A4;X;G;;1;1;;;\nA5;X;G;5;;;;;\n;X;G;5;;;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual((l, [(w["tur"], w["kim"]) for w in uy]),
                             ([{"kod": "A1", "ad": "Ürün", "kategori": "(kategori yok)", "fiyat": D("1250.50"), "eski": D(900), "yeni": D("950.25"), "hedef": D("32.5"), "asgari": D(1000), "adet": D(12)},
                               {"kod": "A5", "ad": "X", "kategori": "G", "fiyat": D(5), "eski": None, "yeni": None, "hedef": None, "asgari": None, "adet": None}],
                              [("Mükerrer ürün kodu", "a1"), ("Ürün okunamadı", "A2"), ("Ürün okunamadı", "A3"), ("Ürün okunamadı", "A4")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--urunler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--marj", "35", "--yuvarla", "0,50", "--azami-artis", "15", "--asgari-degisim", "2", "--indirim"]), 0)
            self.assertEqual(main.main(arg + ["--yontem", "katki", "--marj", "150", "--yuvarla", "0.05"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + x) for x in (["--marj", "100"], ["--marj", "0"], ["--marj", "x"], ["--yuvarla", "0"], ["--azami-artis", "x"], ["--azami-artis", "0"], ["--asgari-degisim=-1"])], [2] * 7)
            self.assertEqual(main.main(["--urunler", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Ürün Kodu;Ürün\nA;1\n", "Ürün Kodu;Mevcut Fiyat\n;1\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)
            with self.assertRaises(SystemExit):
                main.main(arg + ["--yontem", "x"])


if __name__ == "__main__":
    unittest.main()
