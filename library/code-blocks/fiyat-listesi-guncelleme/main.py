"""
Fiyat Listesi Güncelleme — Gerisi Bende kod bloğu
Satış › Satış Destek Uzmanı

Maliyet değişimi ve kâr marjı kurallarına göre yeni fiyat listesini ve değişim raporunu hazırlar.
  Brüt marj %          = (fiyat − maliyet) ÷ fiyat × 100
  Hedef marj           = ürün satırındaki hedef → yoksa --marj → o da yoksa mevcut marj (mevcut fiyat ve eski maliyetten)
  Ham fiyat            = yeni maliyet ÷ (1 − hedef marj ÷ 100)          (--yontem marj, varsayılan)
                       = yeni maliyet × (1 + hedef oran ÷ 100)          (--yontem katki: maliyet üstü kâr)
  Yuvarlama            = ham fiyat, --yuvarla basamağının katına yukarı yuvarlanır (örn. 0,50 → 118,46 → 118,50)
  İndirim yapmama      = hesaplanan fiyat mevcut fiyatın altındaysa mevcut fiyat korunur (--indirim verilmedikçe)
  Artış tavanı         = artış --azami-artis yüzdesini aşarsa fiyat mevcut fiyat × (1 + tavan) ile sınırlanır
  Asgari fiyat         = ürün satırındaki asgari fiyatın altına inilmez (tavandan sonra uygulanır)
  Asgari değişim       = değişim --asgari-degisim yüzdesinin altındaysa fiyat değiştirilmez
Fiyatlar ve maliyetler KDV hariç ve aynı para biriminde olmalıdır. Kod rakip fiyatına, talebe ya da sözleşmeye bakmaz.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
from decimal import ROUND_CEILING, ROUND_FLOOR

SUTUN = {"kod": ("urun kodu", "stok kodu", "kod"), "ad": ("urun", "urun adi", "stok adi", "ad"), "kategori": ("kategori", "grup"), "fiyat": ("mevcut fiyat", "liste fiyati", "fiyat", "satis fiyati"),
         "eski": ("eski maliyet", "onceki maliyet", "mevcut maliyet"), "yeni": ("yeni maliyet", "guncel maliyet", "maliyet"), "hedef": ("hedef marj", "hedef marj %", "hedef oran", "marj"), "asgari": ("asgari fiyat", "taban fiyat", "minimum fiyat"),
         "adet": ("adet", "satis adedi", "miktar")}
Z, YUZ = Decimal(0), Decimal(100)
DURUM_RENK = {"Artış": SARI, "İndirim": MAVI, "Aynı": None}


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("kod", "fiyat")):
        kod = metin(r["kod"])
        if not kod:
            continue
        d = {a: para(r[a]) for a in ("fiyat", "eski", "yeni", "hedef", "asgari", "adet")}
        bozuk = [a for a in d if metin(r[a]) and d[a] is None] + [a for a in ("fiyat", "eski", "yeni") if d[a] is not None and d[a] <= 0] + [a for a in ("asgari", "adet") if d[a] is not None and d[a] < 0]
        if d["fiyat"] is None or bozuk:
            uy.append(uyari("Yüksek", "Ürün okunamadı", kod, f"Satır {r['_satir']}: fiyat ve maliyetler pozitif sayı olmalı (okunamayan: {', '.join(SUTUN[a][0] for a in bozuk) or 'mevcut fiyat'}). Satır alınmadı."))
            continue
        if katla(kod) in gorulen:
            uy.append(uyari("Orta", "Mükerrer ürün kodu", kod, f"Satır {r['_satir']}: aynı kod daha önce geçti; bu satır alınmadı."))
            continue
        gorulen.add(katla(kod))
        l.append({"kod": kod, "ad": metin(r["ad"]), "kategori": metin(r["kategori"]) or "(kategori yok)", **d})
    return l, uy


def _marj(fiyat: Decimal, maliyet: Decimal | None) -> Decimal | None:
    return yuvarla((fiyat - maliyet) * YUZ / fiyat) if maliyet is not None and fiyat else None


def _kat(x: Decimal, basamak: Decimal, yon) -> Decimal:
    return (x / basamak).to_integral_value(rounding=yon) * basamak


def guncelle(urunler: list[dict], marj: Decimal | None = None, yontem: str = "marj", basamak: Decimal = Decimal("0.01"), azami: Decimal | None = None, asgari_degisim: Decimal = Z, indirim: bool = False) -> dict:
    if not urunler:
        raise ValueError("Güncellenecek ürün yok.")
    uy = []
    for u in urunler:
        kim = f"{u['kod']} · {u['ad']}".strip(" ·")
        f, notlar = u["fiyat"], []
        u["eski_marj"], u["degismezse"] = _marj(f, u["eski"]), _marj(f, u["yeni"])
        hedef = u["hedef"] if u["hedef"] is not None else marj
        koru = hedef is None and u["eski"] is not None
        if koru:
            hedef = (f - u["eski"]) * YUZ / (f if yontem == "marj" else u["eski"])
        u["hedef_kullanilan"], yeni = (yuvarla(hedef) if hedef is not None else None), f
        if u["yeni"] is None:
            notlar.append("Yeni maliyet yok")
            uy.append(uyari("Orta", "Yeni maliyet yok", kim, "Yeni maliyet yazılmamış; fiyat değiştirilmedi."))
        elif hedef is None or (yontem == "marj" and hedef >= 100):
            notlar.append("Hedef belirlenemedi")
            uy.append(uyari("Orta", "Hedef marj belirlenemedi", kim, "Hedef marj yok ve eski maliyet yazılmadığı için mevcut marj da bilinmiyor (ya da hedef marj %100 ve üstünde). Fiyat değiştirilmedi."))
        else:
            ham = f * u["yeni"] / u["eski"] if koru else u["yeni"] / (1 - hedef / YUZ) if yontem == "marj" else u["yeni"] * (1 + hedef / YUZ)
            ham = yuvarla(ham, "0.000001")   # bölmeden kalan kırıntı yukarı yuvarlamayı bir basamak atlatmasın
            yeni = _kat(ham, basamak, ROUND_CEILING)
            if yeni < f and not indirim:
                yeni = f
                notlar.append("İndirim yapılmadı")
                uy.append(uyari("Bilgi", "Maliyet düştü, fiyat korundu", kim, f"Hesaplanan fiyat {_kat(ham, basamak, ROUND_CEILING):,.2f}, mevcut fiyat {f:,.2f}. İndirim için --indirim verin."))
            if azami is not None and yeni > f * (1 + azami / YUZ):
                yeni = max(_kat(f * (1 + azami / YUZ), basamak, ROUND_FLOOR), f)
                notlar.append(f"Artış tavanı (%{azami}) uygulandı")
            if u["asgari"] is not None and yeni < u["asgari"]:
                yeni = u["asgari"]
                notlar.append("Asgari fiyat uygulandı")
                uy.append(uyari("Bilgi", "Asgari fiyat uygulandı", kim, f"Hesaplanan fiyat asgari fiyatın ({u['asgari']:,.2f}) altında kaldı; asgari fiyat yazıldı."))
            elif yeni != f and abs(yeni - f) * YUZ / f < asgari_degisim:
                yeni = f
                notlar.append(f"Değişim %{asgari_degisim} eşiğinin altında")
            gercek = _marj(yeni, u["yeni"])
            if yeni < u["yeni"]:
                uy.append(uyari("Yüksek", "Maliyetin altında fiyat", kim, f"Yeni fiyat {yeni:,.2f}, yeni maliyet {u['yeni']:,.2f} (marj %{gercek}). {'; '.join(notlar) or 'Hedef marj eksi.'}"))
            elif yontem == "marj" and gercek < yuvarla(hedef) and yeni != _kat(ham, basamak, ROUND_CEILING):
                uy.append(uyari("Orta", "Hedef marjın altında", kim, f"Hedef marj %{yuvarla(hedef)}, yeni fiyatla marj %{gercek}. {'; '.join(notlar)}."))
        u["yeni_fiyat"], u["nota"] = yeni, "; ".join(notlar)
        u["fark"], u["oran_ham"] = yeni - f, (yeni - f) * YUZ / f
        u["degisim"], u["yeni_marj"] = yuvarla(u["oran_ham"]), _marj(yeni, u["yeni"] if u["yeni"] is not None else u["eski"])
        u["durum"] = "Artış" if yeni > f else "İndirim" if yeni < f else "Aynı"
        u["ciro_etkisi"] = yuvarla(u["fark"] * u["adet"]) if u["adet"] is not None else None

    def ozet(liste: list[dict]) -> dict:
        adetli = [u for u in liste if u["adet"] is not None]
        ciro = sum((u["fiyat"] * u["adet"] for u in adetli), Z)
        etki = sum((u["ciro_etkisi"] for u in adetli), Z)
        return {"urun": len(liste), "artis": sum(1 for u in liste if u["durum"] == "Artış"), "indirim": sum(1 for u in liste if u["durum"] == "İndirim"), "ayni": sum(1 for u in liste if u["durum"] == "Aynı"),
                "ort_degisim": yuvarla(sum((u["oran_ham"] for u in liste), Z) / len(liste)), "ciro_etkisi": etki if adetli else None, "ciro_etkisi_oran": yuvarla(etki * YUZ / ciro) if ciro else None}

    kategoriler = {}
    for u in urunler:
        kategoriler.setdefault(katla(u["kategori"]), []).append(u)
    return {"urunler": urunler, "ozet": ozet(urunler), "kategoriler": [{"kategori": g[0]["kategori"], **ozet(g)} for _, g in sorted(kategoriler.items())], "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Yeni Fiyat Listesi"
    _baslik(ws, ["Ürün Kodu", "Ürün", "Kategori", "Yeni Fiyat (KDV hariç)"], (14, 34, 18, 16))
    for u in r["urunler"]:
        _satir(ws, [u["kod"], u["ad"], u["kategori"], u["yeni_fiyat"]], para_s=(4,))
    ws = wb.create_sheet("Değişim Raporu")
    _baslik(ws, ["Ürün Kodu", "Ürün", "Kategori", "Mevcut Fiyat", "Eski Maliyet", "Yeni Maliyet", "Eski Marj %", "Fiyat Değişmezse Marj %", "Hedef %", "Yeni Fiyat", "Fark", "Değişim %", "Yeni Marj %", "Durum", "Adet", "Ciro Etkisi",
                 "Uygulanan Kural"], (14, 30, 16, 13, 13, 13, 10, 12, 10, 13, 12, 10, 10, 10, 10, 14, 44))
    for u in r["urunler"]:
        _satir(ws, [u["kod"], u["ad"], u["kategori"], u["fiyat"], u["eski"], u["yeni"], u["eski_marj"], u["degismezse"], u["hedef_kullanilan"], u["yeni_fiyat"], u["fark"] or None, u["degisim"], u["yeni_marj"], u["durum"], u["adet"],
                    u["ciro_etkisi"], u["nota"]], para_s=(4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 16),
               dolgu={14: DURUM_RENK[u["durum"]], 13: KIRMIZI if u["yeni_marj"] is not None and u["yeni_marj"] < 0 else None, 8: KIRMIZI if u["degismezse"] is not None and u["degismezse"] < 0 else None}, sar=(17,))
    ws = wb.create_sheet("Özet")
    _baslik(ws, ["Kategori", "Ürün", "Artış", "İndirim", "Aynı", "Ortalama Değişim %", "Ciro Etkisi", "Ciro Etkisi %"], (24, 8, 8, 8, 8, 12, 16, 11))
    for o in r["kategoriler"] + [{"kategori": "TOPLAM", **r["ozet"]}]:
        n = _satir(ws, [o["kategori"], o["urun"], o["artis"], o["indirim"], o["ayni"], o["ort_degisim"], o["ciro_etkisi"], o["ciro_etkisi_oran"]], para_s=(6, 7, 8))
    for h in ws[n]:
        h.font = Font(bold=True)
    ws.append([])
    ws.append(["Ortalama değişim, ürünlerin değişim yüzdelerinin basit ortalamasıdır. Ciro etkisi = (yeni fiyat − mevcut fiyat) × adet; adedin değişmeyeceği varsayılır."])
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, **kw) -> dict:
    urunler, uy = oku(yol)
    r = guncelle(urunler, **kw)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Maliyet değişimi ve marj kurallarına göre yeni fiyat listesini ve değişim raporunu hazırlar.")
    p.add_argument("--urunler", type=Path, help="Ürün Kodu, Ürün, Kategori, Mevcut Fiyat, Eski Maliyet, Yeni Maliyet, Hedef Marj %%, Asgari Fiyat, Adet")
    p.add_argument("--marj", help="Hedefi yazılmamış ürünler için hedef yüzde; verilmezse her ürünün mevcut marjı korunur")
    p.add_argument("--yontem", choices=("marj", "katki"), default="marj", help="marj: satış fiyatı üzerinden brüt marj (varsayılan); katki: maliyet üstü kâr")
    p.add_argument("--yuvarla", help="Fiyatın yukarı yuvarlanacağı basamak (örn. 0,50 veya 5); varsayılan 0,01")
    p.add_argument("--azami-artis", help="Tek seferde izin verilen en yüksek artış yüzdesi (verilmezse tavan yok)")
    p.add_argument("--asgari-degisim", default="0", help="Değişim bu yüzdenin altındaysa fiyat değiştirilmez (varsayılan 0)")
    p.add_argument("--indirim", action="store_true", help="Maliyet düştüğünde fiyatın da inmesine izin ver")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "fiyat_listesi_guncelleme.xlsx")
    a = p.parse_args(argv)
    ornek = a.urunler is None
    yol = ORNEK / "urunler.csv" if ornek else a.urunler
    olcu = lambda s: para(s.replace(".", ",") if s and "," not in s else s)   # noqa: E731
    basamak_s, azami_s = (a.yuvarla or "0,50", a.azami_artis or "20") if ornek else (a.yuvarla or "0,01", a.azami_artis)
    marj, basamak, azami, asgari = olcu(a.marj), olcu(basamak_s), olcu(azami_s), olcu(a.asgari_degisim)
    if (a.marj and (marj is None or marj <= 0 or (a.yontem == "marj" and marj >= 100))) or basamak is None or basamak <= 0 or (azami_s and (azami is None or azami <= 0)) or asgari is None or asgari < 0:
        print("[X] --marj pozitif (marj yönteminde 100'den küçük); --yuvarla ve --azami-artis pozitif; --asgari-degisim sıfır veya pozitif olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, marj=marj, yontem=a.yontem, basamak=basamak, azami=azami, asgari_degisim=asgari, indirim=a.indirim)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {o['urun']} ürün · artış {o['artis']}, indirim {o['indirim']}, aynı {o['ayni']} · ortalama değişim %{o['ort_degisim']}" + (f" · ciro etkisi {o['ciro_etkisi']:,.2f} (%{o['ciro_etkisi_oran']})" if o["ciro_etkisi_oran"] is not None else ""))
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
