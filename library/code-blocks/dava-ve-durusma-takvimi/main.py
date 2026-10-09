"""
Dava ve Duruşma Takvimi — Workers / Workless kod bloğu
Hukuk › Avukat

Dava listesinden duruşma takvimi, usul süreleri listesi ve hatırlatmalar çıkarır:
  - Süre başlangıcı (tebliğ vb.) + olay türüne göre kanuni süre (HMK / İYUK tablosu) veya dosyada yazılı süre
    ("2 hafta", "1 ay", "15 gün", "10").
  - Süre hesabı (HMK 92): gün olarak belirlenen sürede başlangıç günü sayılmaz; hafta / ay olarak belirlenen süre son
    hafta / ayda başlangıç gününe karşılık gelen günde biter, o gün yoksa ayın son günü.
  - Adli tatil (20 Temmuz – 31 Ağustos): adli tatile tabi işlerde son günü tatile rastlayan süre tatilin bittiği günden
    itibaren bir hafta uzar (HMK 104; İYUK 8 → 7 Eylül).
  - Son gün hafta sonu veya resmî tatile rastlarsa izleyen ilk iş günü (HMK 93). Ulusal bayramlar kodda; dinî bayramlar
    tatil dosyasından. Yarım gün (arife, 28 Ekim) son güne denk gelirse uyarı.
  - İç hedef tarih: kanuni son günden --tampon iş günü önce.
  - Duruşmalar: yaklaşanlar, aynı avukatın aynı gün iki duruşması (saatler yakınsa yüksek önem), tarihi geçmiş ama
    güncellenmemiş duruşma, adli tatile rastlayan duruşma.
Rapor: takvim (tarih sırasıyla), süreler (hesap adımlarıyla), duruşmalar, avukat bazında yük, hatırlatmalar, uyarılar
ve takvim uygulamalarına aktarılabilir .ics dosyası. İnternete bağlanmaz.

Kullanım:
    python main.py                                                   # örnek: 12 dava, 12 süre olayı
    python main.py --davalar davalar.xlsx --sureler sureler.xlsx --tatiller tatiller.csv --bugun 09.10.2026 --gun 60
"""
from __future__ import annotations

import argparse
import calendar
import csv
import hashlib
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")

# Olay türü → (usul, süre metni, dayanak). Anahtar katlanmış metinde aranan ifadedir; ilk uyan kullanılır.
SURE_TABLOSU = [
    ("HMK", "cevaba cevap", "2 hafta", "HMK 136"),
    ("HMK", "ikinci cevap", "2 hafta", "HMK 136"),
    ("HMK", "istinafa cevap", "2 hafta", "HMK 347"),
    ("HMK", "cevap", "2 hafta", "HMK 127"),
    ("HMK", "bilirkisi", "2 hafta", "HMK 281"),
    ("HMK", "istinaf", "2 hafta", "HMK 345"),
    ("HMK", "temyiz", "1 ay", "HMK 361"),
    ("IYUK", "cevap", "30 gün", "İYUK 16"),
    ("IYUK", "istinaf", "30 gün", "İYUK 45"),
    ("IYUK", "temyiz", "30 gün", "İYUK 46"),
]
# 2429 sayılı Kanun: sabit tarihli ulusal bayram ve genel tatiller (ay, gün) → ad; yarım günler ayrı
SABIT_TATILLER = {(1, 1): "Yılbaşı", (4, 23): "Ulusal Egemenlik ve Çocuk Bayramı", (5, 1): "Emek ve Dayanışma Günü",
                  (5, 19): "Atatürk'ü Anma, Gençlik ve Spor Bayramı", (7, 15): "Demokrasi ve Millî Birlik Günü", (8, 30): "Zafer Bayramı",
                  (10, 29): "Cumhuriyet Bayramı"}
SABIT_YARIM = {(10, 28): "Cumhuriyet Bayramı arifesi (öğleden sonra)"}
EVET = {"evet", "e", "var", "tamam", "tamamlandi", "yapildi", "x", "1"}
KAPALI_DURUM = {"kapandi", "kesinlesti", "dustu", "feragat", "sulh", "arsiv", "karar kesinlesti"}

DAVA_SUTUNLARI = {"dosya": ("dosya no", "esas no", "dosya"), "mahkeme": ("mahkeme",), "usul": ("usul", "yargi yolu"), "tur": ("dava turu", "konu"),
                  "muvekkil": ("muvekkil",), "karsi": ("karsi taraf",), "sifat": ("muvekkil sifati", "sifat"), "avukat": ("sorumlu avukat", "avukat"),
                  "durusma": ("durusma tarihi", "sonraki durusma"), "saat": ("durusma saati", "saat"), "adli_tatil": ("adli tatile tabi", "adli tatil"),
                  "durum": ("durum",), "not": ("not", "aciklama")}
SURE_SUTUNLARI = {"dosya": ("dosya no", "esas no", "dosya"), "olay": ("olay", "olay turu", "islem"), "baslangic": ("baslangic tarihi", "teblig tarihi", "baslangic"),
                  "sure": ("sure",), "aciklama": ("aciklama", "not"), "tamam": ("tamamlandi", "yapildi", "durum")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def saat(x) -> time | None:
    if isinstance(x, time):
        return x
    if isinstance(x, datetime):
        return x.time()
    m = re.fullmatch(r"(\d{1,2})[:.](\d{2})", str(x or "").strip())
    return time(int(m[1]), int(m[2])) if m and int(m[1]) < 24 and int(m[2]) < 60 else None


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


def usul_bul(x) -> str:
    k = katla(x)
    return "IYUK" if k in {"iyuk", "idari", "idari yargi", "vergi"} or "idari" in k or "iyuk" in k else "HMK"


# ----------------------------------------------------------------------------
# Tatil takvimi ve süre hesabı
# ----------------------------------------------------------------------------

class Takvim:
    def __init__(self, ek_tatiller: dict[date, str] | None = None, ek_yarim: dict[date, str] | None = None):
        self.ek = ek_tatiller or {}
        self.yarim_ek = ek_yarim or {}
        self.yillar = {d.year for d in self.ek} | {d.year for d in self.yarim_ek}

    def tatil_mi(self, d: date) -> str | None:
        if d.weekday() == 5:
            return "Cumartesi"
        if d.weekday() == 6:
            return "Pazar"
        return SABIT_TATILLER.get((d.month, d.day)) or self.ek.get(d)

    def yarim_gun(self, d: date) -> str | None:
        return SABIT_YARIM.get((d.month, d.day)) or self.yarim_ek.get(d)

    def is_gunu_sonra(self, d: date) -> date:
        while self.tatil_mi(d):
            d += timedelta(days=1)
        return d

    def is_gunu_once(self, d: date, n: int) -> date:
        while n > 0:
            d -= timedelta(days=1)
            if not self.tatil_mi(d):
                n -= 1
        return d


def sure_coz(s) -> tuple[int, str] | None:
    """'2 hafta' → (2, 'hafta'); '1 ay' → (1, 'ay'); '15 gün' veya '15' → (15, 'gun')."""
    k = katla(s)
    m = re.fullmatch(r"(\d+)\s*(gun|hafta|ay|yil)?", k)
    if not m:
        return None
    return int(m[1]), (m[2] or "gun")


def ay_ekle(d: date, n: int) -> date:
    y, m = divmod(d.month - 1 + n, 12)
    yil, ay = d.year + y, m + 1
    return date(yil, ay, min(d.day, calendar.monthrange(yil, ay)[1]))


def adli_tatilde(d: date) -> bool:
    return date(d.year, 7, 20) <= d <= date(d.year, 8, 31)


def son_gun(baslangic: date, miktar: int, birim: str, takvim: Takvim, adli_tatil: bool) -> tuple[date, list[str]]:
    adim = []
    if birim == "gun":
        d = baslangic + timedelta(days=miktar)
        adim.append(f"{baslangic:%d.%m.%Y} + {miktar} gün (başlangıç günü sayılmaz) = {d:%d.%m.%Y}")
    elif birim == "hafta":
        d = baslangic + timedelta(weeks=miktar)
        adim.append(f"{baslangic:%d.%m.%Y} + {miktar} hafta (aynı gün) = {d:%d.%m.%Y}")
    elif birim == "ay":
        d = ay_ekle(baslangic, miktar)
        adim.append(f"{baslangic:%d.%m.%Y} + {miktar} ay = {d:%d.%m.%Y}" + (" (karşılık gelen gün yok → ayın son günü)" if d.day != baslangic.day else ""))
    else:
        d = ay_ekle(baslangic, 12 * miktar)
        adim.append(f"{baslangic:%d.%m.%Y} + {miktar} yıl = {d:%d.%m.%Y}")
    if adli_tatil and adli_tatilde(d):
        yeni = date(d.year, 8, 31) + timedelta(days=7)
        adim.append(f"Son gün adli tatilde → tatil bitiminden itibaren 1 hafta: {yeni:%d.%m.%Y}")
        d = yeni
    t = takvim.tatil_mi(d)
    if t:
        yeni = takvim.is_gunu_sonra(d)
        adim.append(f"{d:%d.%m.%Y} {t} → izleyen ilk iş günü {yeni:%d.%m.%Y}")
        d = yeni
    return d, adim


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Dava:
    dosya: str
    mahkeme: str
    usul: str
    tur: str
    muvekkil: str
    karsi: str
    sifat: str
    avukat: str
    durusma: date | None
    saat: time | None
    adli_tatil: bool
    durum: str
    not_: str = ""

    @property
    def acik(self) -> bool:
        return katla(self.durum) not in KAPALI_DURUM


@dataclass
class Sure:
    dosya: str
    olay: str
    baslangic: date
    sure_metni: str
    dayanak: str
    tamam: bool
    aciklama: str
    dava: Dava | None = None
    son: date | None = None
    hedef: date | None = None
    adimlar: list[str] = field(default_factory=list)
    durum: str = ""


def tatilleri_oku(yol: Path | None) -> tuple[dict[date, str], dict[date, str]]:
    tam, yarim = {}, {}
    if not yol:
        return tam, yarim
    for r in tablo_oku(yol)[1:]:
        d = tarih(r[0]) if r else None
        if not d:
            continue
        ad = metin(r[1]) if len(r) > 1 else "Tatil"
        if len(r) > 2 and katla(r[2]) in EVET:
            yarim[d] = ad
        else:
            tam[d] = ad
    return tam, yarim


def davalari_oku(yol: Path) -> list[Dava]:
    davalar = []
    for r in kayitlar(yol, DAVA_SUTUNLARI, ("dosya",)):
        if not metin(r.get("dosya")):
            continue
        at = katla(r.get("adli_tatil"))
        davalar.append(Dava(metin(r["dosya"]), metin(r.get("mahkeme")), usul_bul(r.get("usul")), metin(r.get("tur")), metin(r.get("muvekkil")),
                            metin(r.get("karsi")), metin(r.get("sifat")), metin(r.get("avukat")) or "(atanmamış)", tarih(r.get("durusma")), saat(r.get("saat")),
                            at not in {"hayir", "h", "yok", "0"}, metin(r.get("durum")) or "Derdest", metin(r.get("not"))))
    return davalar


def sureleri_oku(yol: Path | None, davalar: dict[str, Dava]) -> tuple[list[Sure], list[dict]]:
    sureler, uy = [], []
    if not yol:
        return sureler, uy
    for r in kayitlar(yol, SURE_SUTUNLARI, ("dosya", "olay", "baslangic")):
        dosya, olay = metin(r.get("dosya")), metin(r.get("olay"))
        if not dosya or not olay:
            continue
        b = tarih(r.get("baslangic"))
        if not b:
            uy.append({"onem": "Yüksek", "tur": "Başlangıç tarihi yok", "kim": dosya, "aciklama": f"{olay}: başlangıç (tebliğ) tarihi okunamadı; süre hesaplanmadı"})
            continue
        dava = davalar.get(katla(dosya))
        usul = dava.usul if dava else "HMK"
        sure_metni, dayanak = metin(r.get("sure")), ""
        if not sure_metni:
            k = katla(olay)
            uyan = next((x for x in SURE_TABLOSU if x[0] == usul and x[1] in k), None)
            if uyan:
                sure_metni, dayanak = uyan[2], uyan[3]
        else:
            dayanak = "Dosyada belirtilen süre"
        sureler.append(Sure(dosya, olay, b, sure_metni, dayanak, katla(r.get("tamam")) in EVET, metin(r.get("aciklama")), dava))
        if not dava:
            uy.append({"onem": "Orta", "tur": "Dava listesinde yok", "kim": dosya, "aciklama": f"{olay}: dosya dava listesinde bulunamadı; HMK ve adli tatile tabi sayıldı"})
    return sureler, uy


# ----------------------------------------------------------------------------
# Analiz
# ----------------------------------------------------------------------------

def analiz_et(davalar: list[Dava], sureler: list[Sure], takvim: Takvim, bugun: date, gun: int = 60, tampon: int = 3, hatirlatma=(7, 3, 1),
              cakisma_dk: int = 120) -> dict:
    uy = []
    sinir = bugun + timedelta(days=gun)
    for s in sureler:
        c = sure_coz(s.sure_metni) if s.sure_metni else None
        if not c:
            s.durum = "Süre belirsiz"
            uy.append({"onem": "Yüksek", "tur": "Süre belirlenemedi", "kim": s.dosya,
                       "aciklama": f"{s.olay}: olay türü süre tablosunda yok ve 'Süre' sütunu boş/okunamadı ('{s.sure_metni}'). Kesin süreyi yazın"})
            continue
        tabi = s.dava.adli_tatil if s.dava else True
        s.son, s.adimlar = son_gun(s.baslangic, c[0], c[1], takvim, tabi)
        s.hedef = max(min(s.son, takvim.is_gunu_once(s.son, tampon)), min(s.son, bugun))
        if s.son.year not in takvim.yillar:
            s.adimlar.append(f"Not: {s.son.year} dinî bayram tarihleri tatil dosyasında yok")
        yarim = takvim.yarim_gun(s.son)
        if yarim:
            s.adimlar.append(f"Son gün yarım gün: {yarim}")
        kalan = (s.son - bugun).days
        if s.tamam:
            s.durum = "Tamamlandı"
        elif kalan < 0:
            s.durum = "Süre geçti"
            uy.append({"onem": "Yüksek", "tur": "Süre geçmiş", "kim": s.dosya,
                       "aciklama": f"{s.olay}: son gün {s.son:%d.%m.%Y} ({-kalan} gün önce) ve 'tamamlandı' işaretli değil. Hemen kontrol edin"})
        else:
            s.durum = "Bugün son gün" if kalan == 0 else f"{kalan} gün kaldı"
            if kalan <= 7:
                uy.append({"onem": "Yüksek" if kalan <= 3 else "Orta", "tur": "Yaklaşan kesin süre", "kim": s.dosya,
                           "aciklama": f"{s.olay}: son gün {s.son:%d.%m.%Y} ({s.durum}); iç hedef {s.hedef:%d.%m.%Y}"})
        if not s.tamam and yarim:
            uy.append({"onem": "Orta", "tur": "Son gün yarım gün", "kim": s.dosya,
                       "aciklama": f"{s.olay}: son gün {s.son:%d.%m.%Y} {yarim}; işlemi öğleden önce veya daha erken yapın"})
        if not s.tamam and s.son.year not in takvim.yillar:
            uy.append({"onem": "Orta", "tur": "Tatil dosyası eksik", "kim": s.dosya,
                       "aciklama": f"{s.olay}: {s.son.year} yılının dinî bayram tarihleri tatil dosyasında yok; son günü teyit edin"})
        # Sınır durum: ham son gün tatil öncesi hafta sonu, kaydırma adli tatile düşüyor
        if s.dava is None or s.dava.adli_tatil:
            ham = s.baslangic + (timedelta(days=c[0]) if c[1] == "gun" else timedelta(weeks=c[0]) if c[1] == "hafta" else timedelta(0))
            if c[1] in ("gun", "hafta") and not adli_tatilde(ham) and adli_tatilde(takvim.is_gunu_sonra(ham)) and not s.tamam:
                uy.append({"onem": "Orta", "tur": "Adli tatil sınırı", "kim": s.dosya,
                           "aciklama": f"{s.olay}: süre {ham:%d.%m.%Y} tatil gününde bitiyor, izleyen iş günü adli tatilde. Uzamaya güvenmeden "
                                       f"{takvim.is_gunu_once(ham, 1):%d.%m.%Y} tarihine kadar işlem yapın"})

    durusmalar = [d for d in davalar if d.durusma and d.acik]
    for d in davalar:
        if d.acik and d.durusma and d.durusma < bugun:
            uy.append({"onem": "Orta", "tur": "Duruşma tarihi güncellenmemiş", "kim": d.dosya,
                       "aciklama": f"Son duruşma {d.durusma:%d.%m.%Y}; dava açık ama yeni tarih girilmemiş. Tensip / ara kararı kontrol edin"})
        elif d.acik and not d.durusma:
            uy.append({"onem": "Bilgi", "tur": "Duruşma tarihi yok", "kim": d.dosya, "aciklama": f"{d.mahkeme}: duruşma tarihi girilmemiş"})
        if d.acik and d.durusma and d.durusma >= bugun and d.adli_tatil and adli_tatilde(d.durusma):
            uy.append({"onem": "Bilgi", "tur": "Adli tatilde duruşma", "kim": d.dosya,
                       "aciklama": f"{d.durusma:%d.%m.%Y}: adli tatile tabi işte tatil içinde duruşma; tarihi teyit edin (HMK 103 istisnaları hariç)"})
        if d.acik and d.durusma and d.durusma >= bugun and takvim.tatil_mi(d.durusma):
            uy.append({"onem": "Orta", "tur": "Tatil gününde duruşma", "kim": d.dosya, "aciklama": f"{d.durusma:%d.%m.%Y} {takvim.tatil_mi(d.durusma)}; tarihi teyit edin"})
    gelecek = [d for d in durusmalar if d.durusma >= bugun]
    gruplar = defaultdict(list)
    for d in gelecek:
        gruplar[(katla(d.avukat), d.durusma)].append(d)
    for (_, g), lst in gruplar.items():
        if len(lst) < 2:
            continue
        lst.sort(key=lambda d: d.saat or time(0))
        yakin = any(a.saat and b.saat and (datetime.combine(g, b.saat) - datetime.combine(g, a.saat)).seconds / 60 < cakisma_dk for a, b in zip(lst, lst[1:]))
        uy.append({"onem": "Yüksek" if yakin else "Orta", "tur": "Duruşma çakışması", "kim": lst[0].avukat,
                   "aciklama": f"{g:%d.%m.%Y}: " + "; ".join(f"{d.dosya} {d.mahkeme} {d.saat.strftime('%H:%M') if d.saat else 'saat yok'}" for d in lst)
                               + (f" — aralık {cakisma_dk} dakikadan az; yetki belgesi / tevkil düşünün" if yakin else "")})

    # Takvim ve hatırlatmalar
    olaylar = []
    for d in gelecek:
        if d.durusma <= sinir:
            olaylar.append({"tarih": d.durusma, "saat": d.saat, "tur": "Duruşma", "dosya": d.dosya, "mahkeme": d.mahkeme, "avukat": d.avukat,
                            "aciklama": f"{d.tur} — {d.muvekkil} / {d.karsi}".strip(" —/"), "kaynak": d})
    for s in sureler:
        if s.son and not s.tamam and s.son >= bugun and s.son <= sinir:
            olaylar.append({"tarih": s.son, "saat": None, "tur": "Kesin süre son günü", "dosya": s.dosya, "mahkeme": s.dava.mahkeme if s.dava else "",
                            "avukat": s.dava.avukat if s.dava else "", "aciklama": f"{s.olay} ({s.sure_metni}; {s.dayanak}); iç hedef {s.hedef:%d.%m.%Y}", "kaynak": s})
    olaylar.sort(key=lambda o: (o["tarih"], o["saat"] or time(0), o["dosya"]))
    hatirlatmalar = []
    for o in olaylar:
        for n in hatirlatma:
            h = takvim.is_gunu_once(o["tarih"], n)
            if h >= bugun:
                hatirlatmalar.append({"tarih": h, "kac": n, **{k: o[k] for k in ("tur", "dosya", "avukat", "aciklama")}, "olay_tarihi": o["tarih"]})
    hatirlatmalar.sort(key=lambda h: (h["tarih"], h["olay_tarihi"], h["dosya"]))
    yuk = defaultdict(lambda: {"durusma": 0, "sure": 0, "acik_dava": 0})
    for d in davalar:
        if d.acik:
            yuk[d.avukat]["acik_dava"] += 1
    for o in olaylar:
        if o["tarih"] <= bugun + timedelta(days=30):
            yuk[o["avukat"] or "(atanmamış)"]["durusma" if o["tur"] == "Duruşma" else "sure"] += 1
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    uy.sort(key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))
    return {"davalar": davalar, "sureler": sureler, "olaylar": olaylar, "hatirlatmalar": hatirlatmalar, "yuk": dict(yuk), "uyarilar": uy, "bugun": bugun,
            "sinir": sinir, "tampon": tampon}


# ----------------------------------------------------------------------------
# Rapor ve .ics
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
UST = Alignment(vertical="top", wrap_text=True)
GUNLER = ["Pazartesi", "Salı", "Çarşamba", "Perşembe", "Cuma", "Cumartesi", "Pazar"]


def _baslik(ws, basliklar, genislik=()):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def rapor_yaz(cikti: Path, s: dict) -> None:
    wb = Workbook()
    tk = wb.active
    tk.title = "Takvim"
    _baslik(tk, ["Tarih", "Gün", "Saat", "Tür", "Dosya No", "Mahkeme", "Avukat", "Açıklama", "Kalan Gün"], (11, 10, 6, 18, 16, 30, 16, 60, 8))
    for o in s["olaylar"]:
        tk.append([o["tarih"], GUNLER[o["tarih"].weekday()], o["saat"].strftime("%H:%M") if o["saat"] else "", o["tur"], o["dosya"], o["mahkeme"], o["avukat"],
                   o["aciklama"], (o["tarih"] - s["bugun"]).days])
        tk.cell(tk.max_row, 1).number_format = "DD.MM.YYYY"
        if o["tur"] != "Duruşma":
            tk.cell(tk.max_row, 4).fill = PatternFill("solid", fgColor="FDE2E1" if (o["tarih"] - s["bugun"]).days <= 3 else "FFF4CE")
    tk.append([])
    tk.append([f"Rapor tarihi {s['bugun']:%d.%m.%Y}; {s['sinir']:%d.%m.%Y} tarihine kadar olan duruşma ve süre sonları."])

    su = wb.create_sheet("Süreler")
    _baslik(su, ["Dosya No", "Mahkeme", "Olay", "Başlangıç", "Süre", "Dayanak", "Kanuni Son Gün", "Gün", "İç Hedef", "Durum", "Hesap", "Açıklama"],
            (16, 26, 28, 11, 9, 16, 12, 10, 11, 14, 70, 24))
    for x in sorted(s["sureler"], key=lambda x: (x.tamam, x.son or date.max)):
        su.append([x.dosya, x.dava.mahkeme if x.dava else "", x.olay, x.baslangic, x.sure_metni, x.dayanak, x.son, GUNLER[x.son.weekday()] if x.son else "",
                   x.hedef, x.durum, "\n".join(x.adimlar), x.aciklama])
        for j in (4, 7, 9):
            su.cell(su.max_row, j).number_format = "DD.MM.YYYY"
        su.cell(su.max_row, 11).alignment = UST
        renk = "E3F4E1" if x.tamam else "FDE2E1" if x.durum in ("Süre geçti", "Süre belirsiz") or (x.son and (x.son - s["bugun"]).days <= 3) else None
        if renk:
            su.cell(su.max_row, 10).fill = PatternFill("solid", fgColor=renk)

    du = wb.create_sheet("Duruşmalar")
    _baslik(du, ["Duruşma", "Gün", "Saat", "Dosya No", "Mahkeme", "Usul", "Dava Türü", "Müvekkil", "Sıfat", "Karşı Taraf", "Avukat", "Durum", "Not"],
            (11, 10, 6, 16, 30, 6, 20, 22, 9, 22, 16, 10, 24))
    for d in sorted(s["davalar"], key=lambda d: (d.durusma or date.max, d.saat or time(0))):
        du.append([d.durusma, GUNLER[d.durusma.weekday()] if d.durusma else "", d.saat.strftime("%H:%M") if d.saat else "", d.dosya, d.mahkeme,
                   "İYUK" if d.usul == "IYUK" else "HMK", d.tur, d.muvekkil, d.sifat, d.karsi, d.avukat, d.durum, d.not_])
        du.cell(du.max_row, 1).number_format = "DD.MM.YYYY"
        if d.acik and d.durusma and d.durusma < s["bugun"]:
            du.cell(du.max_row, 1).fill = PatternFill("solid", fgColor="FFF4CE")
    du.auto_filter.ref = f"A1:M{du.max_row}"

    ay = wb.create_sheet("Avukat Bazında")
    _baslik(ay, ["Avukat", "Açık Dava", "30 Gün İçinde Duruşma", "30 Gün İçinde Süre Sonu"], (20, 10, 14, 14))
    for a, v in sorted(s["yuk"].items()):
        ay.append([a, v["acik_dava"], v["durusma"], v["sure"]])

    ha = wb.create_sheet("Hatırlatmalar")
    _baslik(ha, ["Hatırlatma Tarihi", "Kaç İş Günü Önce", "Olay Tarihi", "Tür", "Dosya No", "Avukat", "Açıklama"], (12, 9, 11, 18, 16, 16, 70))
    for h in s["hatirlatmalar"]:
        ha.append([h["tarih"], h["kac"], h["olay_tarihi"], h["tur"], h["dosya"], h["avukat"], h["aciklama"]])
        ha.cell(ha.max_row, 1).number_format = ha.cell(ha.max_row, 3).number_format = "DD.MM.YYYY"

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Dosya / Kim", "Açıklama"], (9, 26, 18, 100))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["kim"], u["aciklama"]])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def _ics_kacis(x: str) -> str:
    return x.replace("\\", "\\\\").replace(";", "\\;").replace(",", "\\,").replace("\n", "\\n")


def _ics_katla(satir: str) -> list[str]:
    """RFC 5545: 75 sekizliyi aşan satırlar boşlukla devam ettirilir (UTF-8 karakter bölünmeden)."""
    parcalar, mevcut = [], ""
    for ch in satir:
        if len((mevcut + ch).encode("utf-8")) > (75 if not parcalar else 74):
            parcalar.append(mevcut)
            mevcut = ""
        mevcut += ch
    parcalar.append(mevcut)
    return [parcalar[0]] + [" " + p for p in parcalar[1:]]


def ics_yaz(yol: Path, s: dict, sure_dk: int = 60, alarm=(7, 1)) -> None:
    damga = s["bugun"].strftime("%Y%m%dT000000Z")
    satirlar = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//Workers Workless//Dava ve Durusma Takvimi//TR", "CALSCALE:GREGORIAN"]
    for o in s["olaylar"]:
        uid = hashlib.sha1(f"{o['tur']}|{o['dosya']}|{o['tarih']}|{o['aciklama']}".encode()).hexdigest()[:20] + "@workers-workless"
        satirlar += ["BEGIN:VEVENT", f"UID:{uid}", f"DTSTAMP:{damga}"]
        if o["saat"]:
            bas = datetime.combine(o["tarih"], o["saat"])
            # Yerel ("floating") saat: takvim uygulaması kullanıcının saat diliminde gösterir
            satirlar += [f"DTSTART:{bas:%Y%m%dT%H%M%S}", f"DTEND:{bas + timedelta(minutes=sure_dk):%Y%m%dT%H%M%S}"]
        else:
            satirlar += [f"DTSTART;VALUE=DATE:{o['tarih']:%Y%m%d}", f"DTEND;VALUE=DATE:{o['tarih'] + timedelta(days=1):%Y%m%d}"]
        ozet = f"{'Duruşma' if o['tur'] == 'Duruşma' else 'SON GÜN'}: {o['dosya']} {o['mahkeme']}".strip()
        satirlar += [f"SUMMARY:{_ics_kacis(ozet)}", f"DESCRIPTION:{_ics_kacis(o['aciklama'] + (' · Avukat: ' + o['avukat'] if o['avukat'] else ''))}"]
        if o["mahkeme"]:
            satirlar.append(f"LOCATION:{_ics_kacis(o['mahkeme'])}")
        for g in alarm:
            satirlar += ["BEGIN:VALARM", "ACTION:DISPLAY", f"DESCRIPTION:{_ics_kacis(ozet)}", f"TRIGGER:-P{g}D", "END:VALARM"]
        satirlar.append("END:VEVENT")
    satirlar.append("END:VCALENDAR")
    cikis = []
    for x in satirlar:
        cikis += _ics_katla(x)
    yol.write_bytes(("\r\n".join(cikis) + "\r\n").encode("utf-8"))


def calistir(dava_yolu: Path, cikti: Path, bugun: date, sure_yolu: Path | None = None, tatil_yolu: Path | None = None, **kw) -> dict:
    davalar = davalari_oku(dava_yolu)
    if not davalar:
        raise ValueError(f"{dava_yolu.name}: dava bulunamadı")
    sozluk = {katla(d.dosya): d for d in davalar}
    sureler, uy = sureleri_oku(sure_yolu, sozluk)
    takvim = Takvim(*tatilleri_oku(tatil_yolu))
    s = analiz_et(davalar, sureler, takvim, bugun, **kw)
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    s["uyarilar"] = sorted(uy + s["uyarilar"], key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))
    rapor_yaz(cikti, s)
    s["ics"] = cikti.with_suffix(".ics")
    ics_yaz(s["ics"], s)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Dava listesinden duruşma takvimi, kesin süre listesi ve hatırlatmalar çıkarır.")
    p.add_argument("--davalar", type=Path, default=ORNEK / "davalar.csv",
                   help="Dosya No, Mahkeme, Usul (HMK/İYUK), Dava Türü, Müvekkil, Karşı Taraf, Müvekkil Sıfatı, Sorumlu Avukat, Duruşma Tarihi, "
                        "Duruşma Saati, Adli Tatile Tabi (Evet/Hayır), Durum, Not")
    p.add_argument("--sureler", type=Path, help="Dosya No, Olay, Başlangıç Tarihi (tebliğ), Süre (boşsa tablodan), Açıklama, Tamamlandı")
    p.add_argument("--tatiller", type=Path, help="Tarih, Açıklama, Yarım Gün (Evet) — dinî bayramlar ve arifeler")
    p.add_argument("--bugun", help="Rapor tarihi GG.AA.YYYY (örnek veride 09.10.2026)")
    p.add_argument("--gun", type=int, default=60, help="Takvimde kaç gün ileri gösterilsin (varsayılan 60)")
    p.add_argument("--tampon", type=int, default=3, help="İç hedef tarih: son günden kaç iş günü önce (varsayılan 3)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "dava_takvimi.xlsx")
    a = p.parse_args(argv)
    ornek = a.davalar == ORNEK / "davalar.csv"
    sure = a.sureler or (ORNEK / "sureler.csv" if ornek else None)
    tatil = a.tatiller or (ORNEK / "tatiller.csv" if ornek else None)
    for y in (a.davalar, sure, tatil):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    bugun = tarih(a.bugun) if a.bugun else (date(2026, 10, 9) if ornek else date.today())
    if not bugun:
        print("[X] --bugun GG.AA.YYYY biçiminde olmalı")
        return 2
    try:
        s = calistir(a.davalar, a.cikti, bugun, sure, tatil, gun=a.gun, tampon=a.tampon)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {sum(d.acik for d in s['davalar'])} açık dava · {len(s['sureler'])} süre olayı · {bugun:%d.%m.%Y}–{s['sinir']:%d.%m.%Y} arası {len(s['olaylar'])} takvim kaydı")
    for o in s["olaylar"][:8]:
        print(f"     {o['tarih']:%d.%m.%Y} {o['saat'].strftime('%H:%M') if o['saat'] else '     '}  {o['tur']:<20} {o['dosya']}  {o['aciklama'][:60]}")
    for u in s["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['kim']} · {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[OK] Takvim dosyası (.ics): {s['ics'].resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
