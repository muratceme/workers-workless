import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
T = main.Takvim({date(2026, 3, 20): "Ramazan Bayramı 1. gün", date(2026, 3, 21): "Ramazan Bayramı 2. gün", date(2026, 3, 22): "Ramazan Bayramı 3. gün"},
                {date(2026, 3, 19): "Ramazan Bayramı arifesi"})


class SureHesabiTesti(unittest.TestCase):
    def test_hafta_ay_gun(self):
        self.assertEqual(main.son_gun(date(2026, 9, 30), 2, "hafta", T, True)[0], date(2026, 10, 14))
        self.assertEqual(main.son_gun(date(2026, 8, 31), 1, "ay", T, True)[0], date(2026, 9, 30))       # 31 Eylül yok → ayın son günü
        self.assertEqual(main.son_gun(date(2026, 1, 31), 1, "ay", T, True)[0], date(2026, 3, 2))        # 28 Şubat Cumartesi → Pazartesi
        self.assertEqual(main.son_gun(date(2026, 9, 18), 30, "gun", T, True)[0], date(2026, 10, 19))    # Pazar → Pazartesi

    def test_adli_tatil(self):
        d, adim = main.son_gun(date(2026, 7, 24), 2, "hafta", T, True)
        self.assertEqual(d, date(2026, 9, 7))                                    # 07.08 adli tatilde → 31.08 + 1 hafta
        self.assertIn("adli tatil", adim[1])
        self.assertEqual(main.son_gun(date(2026, 7, 24), 2, "hafta", T, False)[0], date(2026, 8, 7))    # tabi olmayan iş
        self.assertEqual(main.son_gun(date(2026, 8, 20), 2, "hafta", T, True)[0], date(2026, 9, 3))     # son gün tatil sonrası: uzama yok
        self.assertEqual(main.son_gun(date(2027, 7, 27), 2, "hafta", T, True)[0], date(2027, 9, 7))     # 07.09.2027 Salı

    def test_bayram_ve_ulusal_tatil(self):
        self.assertEqual(main.son_gun(date(2026, 3, 6), 2, "hafta", T, True)[0], date(2026, 3, 23))     # 20.03 bayram → 23.03
        self.assertEqual(main.son_gun(date(2026, 10, 15), 2, "hafta", T, True)[0], date(2026, 10, 30))  # 29.10 Cumhuriyet Bayramı
        self.assertEqual(T.yarim_gun(date(2026, 10, 28))[:10], "Cumhuriyet")
        self.assertEqual(main.sure_coz("2 Hafta"), (2, "hafta"))
        self.assertEqual(main.sure_coz("15"), (15, "gun"))
        self.assertIsNone(main.sure_coz("makul süre"))


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "d.xlsx"
        cls.s = main.calistir(ORNEK / "davalar.csv", cls.cikti, date(2026, 10, 9), ORNEK / "sureler.csv", ORNEK / "tatiller.csv")
        cls.u = {(u["tur"], u["kim"]) for u in cls.s["uyarilar"]}
        cls.su = {(x.dosya, x.olay): x for x in cls.s["sureler"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_sureler(self):
        self.assertEqual(self.su[("2026/2210 E.", "Savunmaya cevap")].dayanak, "İYUK 16")
        self.assertEqual(self.su[("2026/1902 E.", "İstinaf başvurusu")].son, date(2026, 11, 4))          # İYUK 45: 30 gün
        self.assertEqual(self.su[("2026/301 E.", "Cevaba cevap dilekçesi")].dayanak, "HMK 136")
        self.assertEqual(self.su[("2024/655 E.", "İstinaf başvurusu")].durum, "Tamamlandı")
        self.assertEqual(self.su[("2026/118 E.", "Cevap dilekçesi")].hedef, date(2026, 10, 9))         # 3 iş günü önce
        self.assertEqual(self.su[("2026/520 E.", "Kesin süre - tanık listesi")].durum, "Süre belirsiz")

    def test_uyarilar(self):
        self.assertTrue({("Süre geçmiş", "2023/210 E."), ("Yaklaşan kesin süre", "2026/77 E."), ("Süre belirlenemedi", "2026/520 E."),
                         ("Duruşma çakışması", "Av. Örnek Avukat 1"), ("Duruşma tarihi güncellenmemiş", "2025/980 E."), ("Son gün yarım gün", "2026/640 E."),
                         ("Tatil gününde duruşma", "2026/88 E."), ("Dava listesinde yok", "2023/210 E.")} <= self.u)
        self.assertNotIn(("Duruşma tarihi güncellenmemiş", "2025/1440 E."), self.u)        # kesinleşmiş dosya
        cak = next(u for u in self.s["uyarilar"] if u["tur"] == "Duruşma çakışması" and u["kim"] == "Av. Örnek Avukat 1")
        self.assertEqual(cak["onem"], "Yüksek")
        self.assertEqual(self.s["uyarilar"][0]["onem"], "Yüksek")

    def test_takvim_ve_ics(self):
        o = self.s["olaylar"]
        self.assertEqual([x["tarih"] for x in o], sorted(x["tarih"] for x in o))
        self.assertNotIn("2026/640 E.", {x["dosya"] for x in o if x["tur"] == "Duruşma"})            # 15.12: 60 gün dışında
        ics = self.s["ics"].read_bytes().decode("utf-8")
        self.assertEqual(ics.count("BEGIN:VEVENT"), len(o))
        self.assertIn("DTSTART:20261022T103000", ics)
        self.assertTrue(all(len(x.encode("utf-8")) <= 75 for x in ics.split("\r\n")))
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Takvim", "Süreler", "Duruşmalar", "Avukat Bazında", "Hatırlatmalar", "Uyarılar"])

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--bugun", "x"]), 2)
            self.assertEqual(main.main(["--davalar", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
