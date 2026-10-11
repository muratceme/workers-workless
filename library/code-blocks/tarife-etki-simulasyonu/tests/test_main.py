import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
SEG = ["Sınıf", "Bölge"]


def pol(no, prim, hasar, sinif="A", bolge="X"):
    return {"no": no, "prim": D(prim), "hasar": D(hasar), "seg": {"Sınıf": sinif, "Bölge": bolge}}


def kural(degisim, alan=None, deger="", kayip=0):
    return {"alan": alan, "deger": main.katla(deger), "deger_ad": deger, "degisim": D(degisim), "kayip": D(kayip), "aciklama": ""}


PORTFOY = [pol("P1", 1000, 200, "A"), pol("P2", 1000, 900, "C"), pol("P3", 2000, 1500, "C", "Y")]          # prim 4.000, hasar 2.600, H/P %65


class SimulasyonTesti(unittest.TestCase):
    def test_tek_kural(self):
        r = main.simule(PORTFOY, {"S": [kural(20, "Sınıf", "c")]}, SEG)
        s = r["senaryolar"][0]
        self.assertEqual((r["mevcut"]["prim"], r["mevcut"]["hp"]), (D(4000), D("65.00")))
        self.assertEqual((s["prim"], s["prim_degisim"], s["hp"], s["hp_fark"], s["etkilenen"], s["en_yuksek"], s["en_dusuk"]),
                         (D("4600.00"), D("15.00"), D("56.52"), D("-8.48"), 2, D("20.00"), D("0.00")))                    # 2.600 ÷ 4.600
        self.assertEqual([x["yeni"] for x in s["detay"]], [D("1000.00"), D("1200.00"), D("2400.00")])

    def test_zincirleme_ve_tavan(self):
        k = [kural(10), kural(-5, "Sınıf", "A"), kural(20, "Sınıf", "C")]
        s = main.simule(PORTFOY, {"S": k}, SEG)["senaryolar"][0]
        self.assertEqual([x["yeni"] for x in s["detay"]], [D("1045.00"), D("1320.00"), D("2640.00")])                     # 1,10 × 0,95 ve 1,10 × 1,20
        s = main.simule(PORTFOY, {"S": k}, SEG, azami_artis=D(25))["senaryolar"][0]
        self.assertEqual(([x["yeni"] for x in s["detay"]], s["tavan"], s["en_yuksek"]), ([D("1045.00"), D("1250.00"), D("2500.00")], 2, D("25.00")))

    def test_kayip_varsayimi(self):
        s = main.simule(PORTFOY, {"S": [kural(30, "Sınıf", "C", 15)]}, SEG)["senaryolar"][0]
        # C: prim 3.000 × 1,30 × 0,85 = 3.315; hasar 2.400 × 0,85 = 2.040.  Toplam prim 4.315, hasar 2.240
        self.assertEqual((s["tarife_prim"], s["tarife_degisim"], s["prim"], s["prim_degisim"], s["hasar"], s["hp"], s["kalma"]),
                         (D("4900.00"), D("22.50"), D("4315.00"), D("7.88"), D("2240.00"), D("51.91"), D("88.75")))

    def test_segment_ve_hedef(self):
        r = main.simule(PORTFOY, {"S": [kural(20, "Sınıf", "C")]}, SEG, hedef_hp=D(50))
        d = {(x["alan"], x["deger"]): (x["police"], x["prim"], x["yeni"], x["degisim"], x["hp"], x["yeni_hp"]) for x in r["segment"]}
        self.assertEqual(d[("Sınıf", "C")], (2, D(3000), D("3600.00"), D("20.00"), D("80.00"), D("66.67")))
        self.assertEqual(d[("Sınıf", "A")], (1, D(1000), D("1000.00"), D("0.00"), D("20.00"), D("20.00")))
        self.assertEqual(d[("Bölge", "Y")][3:], (D("20.00"), D("75.00"), D("62.50")))
        self.assertEqual(r["gerekli"], D("30.00"))                                                                         # 65 ÷ 50 − 1
        r = main.simule(PORTFOY, {"S": [kural(20, "Sınıf", "Z")]}, SEG)
        self.assertEqual(([u["tur"] for u in r["uyarilar"]], r["senaryolar"][0]["prim_degisim"], r["gerekli"]), (["Senaryo hiçbir poliçeyi etkilemiyor"], D("0.00"), None))


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "portfoy.csv", ORNEK / "senaryolar.csv", cikti, hedef_hp=D(60))
            m = r["mevcut"]
            self.assertEqual((m["police"], m["prim"], m["hasar"], m["hp"], r["gerekli"]), (12, D(900000), D(595000), D("66.11"), D("10.19")))
            self.assertEqual([(s["prim"], s["prim_degisim"], s["hp"], s["etkilenen"]) for s in r["senaryolar"]],
                             [(D("968000.00"), D("7.56"), D("61.47"), 4), (D("974050.00"), D("8.23"), D("61.09"), 12), (D("935700.00"), D("3.97"), D("57.98"), 4)])
            self.assertEqual((r["senaryolar"][2]["hasar"], r["senaryolar"][2]["tarife_prim"]), (D("542500.00"), D("1002000.00")))
            self.assertEqual(r["uyarilar"], [])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Senaryo Özeti", "Segment Etkisi", "Poliçe Etkisi", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            p, s = Path(tmp) / "p.csv", Path(tmp) / "s.csv"
            p.write_text("Poliçe No;Sınıf;Prim;Hasar\nP1;A;1.000,00;200\nP2;;0;5\nP3;C;500;çok\n;C;1;1\n", encoding="utf-8")
            l, seg, uy = main.portfoy_oku(p)
            self.assertEqual(([(x["no"], x["prim"], x["hasar"], x["seg"]) for x in l], seg), ([("P1", D("1000.00"), D(200), {"Sınıf": "A"}), ("P3", D(500), D(0), {"Sınıf": "C"})], ["Sınıf"]))
            self.assertEqual(sorted(u["tur"] for u in uy), ["Hasar okunamadı", "Prim okunamadı"])
            s.write_text("Senaryo;Alan;Değer;Değişim %;Kayıp %\nS1;sınıf;C;%20;150\nS1;Bölge;X;10;\nS1;*;;çok;\nS2;;;-100;\nS3;Genel;;5;10\n", encoding="utf-8")
            sen, uy = main.senaryo_oku(s, seg)
            self.assertEqual({k: [(x["alan"], x["degisim"], x["kayip"]) for x in v] for k, v in sen.items()}, {"S1": [("Sınıf", D(20), D(0))], "S3": [(None, D(5), D(10))]})
            self.assertEqual(sorted(u["tur"] for u in uy), ["Değişim oranı okunamadı", "Değişim oranı okunamadı", "Kayıp oranı okunamadı", "Kural alanı portföyde yok"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--portfoy", str(p), "--senaryolar", str(s), "--cikti", c]
            self.assertEqual(main.main(arg + ["--azami-artis", "15", "--hedef-hp", "55"]), 0)
            self.assertEqual(main.main(arg + ["--hedef-hp", "0"]), 2)
            self.assertEqual(main.main(["--portfoy", str(Path(tmp) / "yok.csv")]), 1)
            p.write_text("Poliçe No;Prim\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
