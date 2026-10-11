"""
Tarife Etki Simülasyonu — Workers / Workless kod bloğu
Aktüerya › Aktüerya Uzmanı

Tarife değişikliği senaryolarını mevcut portföye uygular; prim ve hasar/prim (H/P) oranına etkisini hesaplar.
  Kural         = Senaryo × Alan × Değer × Değişim % × Kayıp %      (Alan = portföy dosyasındaki bir sütun; * = tüm poliçeler)
  Yeni prim     = prim × Π(1 + değişim)          bir poliçeye uyan tüm kurallar zincirleme uygulanır; --azami-artis ile sınırlanabilir
  Kalma oranı   = Π(1 − kayıp)                   zam nedeniyle ayrılacağı varsayılan pay; prim ve hasar aynı oranda azalır
  Beklenen prim = yeni prim × kalma oranı;   beklenen hasar = hasar × kalma oranı;   H/P = beklenen hasar ÷ beklenen prim
  Hedef H/P için gerekli genel değişim = mevcut H/P ÷ hedef H/P − 1
Statik bir "olduğu gibi" (as-if) hesaptır: hasar eğilimi, enflasyon, IBNR, masraflar ve yeni iş hesaba girmez. Kayıp oranları sizin varsayımınızdır.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SABIT = {"no": ("police no", "police numarasi"), "prim": ("prim", "mevcut prim", "yillik prim", "net prim"), "hasar": ("hasar", "gerceklesen hasar", "hasar tutari")}
SENARYO_SUTUN = {"senaryo": ("senaryo", "senaryo adi"), "alan": ("alan", "sutun"), "deger": ("deger", "segment"), "degisim": ("degisim", "degisim yuzde", "zam", "oran"),
                 "kayip": ("kayip", "kayip yuzde", "beklenen kayip", "ayrilma"), "aciklama": ("aciklama", "not")}
Z, BIR = Decimal(0), Decimal(1)


def portfoy_oku(yol: Path) -> tuple[list[dict], list[str], list[dict]]:
    """(poliçeler, segment sütun adları, uyarılar). Poliçe No, Prim, Hasar dışındaki tüm sütunlar segment sayılır."""
    s, uy = tablo_oku(yol), []
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {k: next((b.index(a) for a in adlar if a in b), None) for k, adlar in SABIT.items()}
        if es["prim"] is not None and es["no"] is not None:
            break
    else:
        raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: Poliçe No, Prim")
    segment = [(j, metin(c)) for j, c in enumerate(s[bi]) if metin(c) and j not in es.values()]
    l = []
    for n, r in enumerate(s[bi + 1:], bi + 2):
        al = lambda j: r[j] if j is not None and j < len(r) else None   # noqa: E731
        no = metin(al(es["no"]))
        if not no:
            continue
        prim, hasar = para(al(es["prim"])), para(al(es["hasar"]))
        if prim is None or prim <= 0:
            uy.append(uyari("Yüksek", "Prim okunamadı", f"{no} (satır {n})", f"'{metin(al(es['prim']))}'. Poliçe simülasyona alınmadı."))
            continue
        if hasar is None and metin(al(es["hasar"])):
            uy.append(uyari("Yüksek", "Hasar okunamadı", f"{no} (satır {n})", f"'{metin(al(es['hasar']))}'. Hasar 0 sayıldı."))
        l.append({"no": no, "prim": prim, "hasar": hasar or Z, "seg": {ad: metin(al(j)) or "Belirtilmemiş" for j, ad in segment}})
    if es["hasar"] is None:
        uy.append(uyari("Orta", "Hasar sütunu yok", yol.name, "Hasar sütunu bulunamadı; H/P oranları 0 görünür, yalnız prim etkisi anlamlıdır."))
    return l, [ad for _, ad in segment], uy


def senaryo_oku(yol: Path, segmentler: list[str]) -> tuple[dict, list[dict]]:
    sen, uy = {}, []
    adlar = {katla(a): a for a in segmentler}
    for r in kayitlar(yol, SENARYO_SUTUN, ("senaryo", "degisim")):
        ad = metin(r["senaryo"])
        if not ad:
            continue
        kim = f"{ad} (satır {r['_satir']})"
        d, k = para(r["degisim"]), para(r["kayip"])
        if d is None or d <= -100:
            uy.append(uyari("Yüksek", "Değişim oranı okunamadı", kim, f"'{metin(r['degisim'])}'. −100'den büyük bir yüzde olmalı. Kural uygulanmadı."))
            continue
        if metin(r["kayip"]) and (k is None or not 0 <= k <= 100):
            uy.append(uyari("Orta", "Kayıp oranı okunamadı", kim, f"'{metin(r['kayip'])}'. 0-100 arası olmalı; kayıp 0 sayıldı."))
            k = None
        alan = metin(r["alan"])
        hepsi = not alan or "*" in alan or katla(alan) in ("hepsi", "tumu", "genel")
        if not hepsi and katla(alan) not in adlar:
            uy.append(uyari("Yüksek", "Kural alanı portföyde yok", kim, f"'{alan}' adında sütun yok (sütunlar: {', '.join(segmentler)}). Kural uygulanmadı."))
            continue
        sen.setdefault(ad, []).append({"alan": None if hepsi else adlar[katla(alan)], "deger": katla(r["deger"]), "deger_ad": metin(r["deger"]), "degisim": d, "kayip": k or Z,
                                       "aciklama": metin(r["aciklama"])})
    if not sen:
        raise ValueError(f"{yol.name}: uygulanabilir senaryo kuralı yok.")
    return sen, uy


def hp(hasar: Decimal, prim: Decimal) -> Decimal | None:
    return yuvarla(hasar / prim * 100) if prim > 0 else None


def degisim(yeni: Decimal, eski: Decimal) -> Decimal | None:
    return yuvarla((yeni - eski) / eski * 100) if eski else None


def uygula(policeler: list[dict], kurallar: list[dict], azami_artis: Decimal | None = None) -> list[dict]:
    """Her poliçe için (çarpan, kalma oranı, uyan kural sayısı, tavan uygulandı mı)."""
    sonuc = []
    for p in policeler:
        c, kalma, n = BIR, BIR, 0
        for k in kurallar:
            if k["alan"] is None or katla(p["seg"].get(k["alan"])) == k["deger"]:
                c, kalma, n = c * (1 + k["degisim"] / 100), kalma * (1 - k["kayip"] / 100), n + 1
        tavan = azami_artis is not None and c > 1 + azami_artis / 100
        if tavan:
            c = 1 + azami_artis / 100
        sonuc.append({"carpan": c, "kalma": kalma, "kural": n, "tavan": tavan, "yeni": yuvarla(p["prim"] * c), "b_prim": p["prim"] * c * kalma, "b_hasar": p["hasar"] * kalma})
    return sonuc


def simule(policeler: list[dict], senaryolar: dict, segmentler: list[str], azami_artis: Decimal | None = None, hedef_hp: Decimal | None = None) -> dict:
    prim, hasar = sum((p["prim"] for p in policeler), Z), sum((p["hasar"] for p in policeler), Z)
    r = {"mevcut": {"police": len(policeler), "prim": prim, "hasar": hasar, "hp": hp(hasar, prim)}, "senaryolar": [], "segment": [], "policeler": policeler, "uyarilar": [],
         "gerekli": yuvarla((hasar / prim * 100 / hedef_hp - 1) * 100) if hedef_hp and prim else None, "hedef": hedef_hp}
    for ad, kurallar in senaryolar.items():
        s = uygula(policeler, kurallar, azami_artis)
        bp, bh = sum((x["b_prim"] for x in s), Z), sum((x["b_hasar"] for x in s), Z)
        tarife = sum((x["yeni"] for x in s), Z)
        artis = [degisim(x["yeni"], p["prim"]) for x, p in zip(s, policeler)]
        ozet = {"ad": ad, "kural": len(kurallar), "etkilenen": sum(x["kural"] > 0 and x["carpan"] != 1 for x in s), "tarife_prim": tarife, "tarife_degisim": degisim(tarife, prim),
                "prim": yuvarla(bp), "prim_degisim": degisim(bp, prim), "hasar": yuvarla(bh), "hp": hp(bh, bp), "tavan": sum(x["tavan"] for x in s), "en_yuksek": max(artis), "en_dusuk": min(artis),
                "kalma": yuvarla(sum((p["prim"] * x["kalma"] for x, p in zip(s, policeler)), Z) / prim * 100) if prim else None, "detay": s,
                "aciklama": "; ".join(f"{k['alan'] or 'Tümü'}{'=' + k['deger_ad'] if k['alan'] else ''} {'+' if k['degisim'] >= 0 else ''}{k['degisim'].normalize():f}%"
                                      + (f" (kayıp %{k['kayip'].normalize():f})" if k["kayip"] else "") for k in kurallar)}
        ozet["hp_fark"] = ozet["hp"] - r["mevcut"]["hp"] if ozet["hp"] is not None and r["mevcut"]["hp"] is not None else None
        if not ozet["etkilenen"]:
            r["uyarilar"].append(uyari("Orta", "Senaryo hiçbir poliçeyi etkilemiyor", ad, "Kurallardaki değerler portföydeki değerlerle eşleşmedi; yazımı kontrol edin."))
        r["senaryolar"].append(ozet)
        for seg in segmentler:
            g = {}
            for p, x in zip(policeler, s):
                y = g.setdefault(p["seg"][seg], {"police": 0, "prim": Z, "hasar": Z, "bp": Z, "bh": Z})
                y["police"] += 1
                y["prim"] += p["prim"]
                y["hasar"] += p["hasar"]
                y["bp"] += x["b_prim"]
                y["bh"] += x["b_hasar"]
            for deger, y in sorted(g.items()):
                r["segment"].append({"senaryo": ad, "alan": seg, "deger": deger, "police": y["police"], "prim": y["prim"], "yeni": yuvarla(y["bp"]), "degisim": degisim(y["bp"], y["prim"]),
                                     "hp": hp(y["hasar"], y["prim"]), "yeni_hp": hp(y["bh"], y["bp"])})
    return r


def rapor_yaz(cikti: Path, r: dict, segmentler: list[str]) -> None:
    wb, m = Workbook(), r["mevcut"]
    ws = wb.active
    ws.title = "Senaryo Özeti"
    _baslik(ws, ["Senaryo", "Kurallar", "Etkilenen Poliçe", "Tarife Primi (kayıpsız)", "Tarife Değişimi %", "Beklenen Prim", "Prim Değişimi %", "Beklenen Hasar", "H/P %", "H/P Farkı (puan)",
                 "Primin Kalma Oranı %", "En Yüksek Artış %", "En Düşük Değişim %", "Tavana Takılan"], (26, 60, 10, 18, 11, 18, 11, 18, 9, 10, 11, 10, 10, 9))
    _satir(ws, ["Mevcut durum", "", m["police"], m["prim"], None, m["prim"], None, m["hasar"], m["hp"], None, 100, None, None, None], para_s=(4, 6, 8, 9))
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for s in r["senaryolar"]:
        _satir(ws, [s["ad"], s["aciklama"], s["etkilenen"], s["tarife_prim"], s["tarife_degisim"], s["prim"], s["prim_degisim"], s["hasar"], s["hp"], s["hp_fark"], s["kalma"], s["en_yuksek"],
                    s["en_dusuk"], s["tavan"] or None], para_s=(4, 5, 6, 7, 8, 9, 10, 11, 12, 13), sar=(2,),
               dolgu={10: None if s["hp_fark"] is None else YESIL if s["hp_fark"] < 0 else KIRMIZI if s["hp_fark"] > 0 else None})
    ws.append([])
    if r["gerekli"] is not None:
        ws.append([f"Hedef H/P %{r['hedef']} için gerekli genel prim değişimi: %{r['gerekli']} (mevcut H/P ÷ hedef H/P − 1; kayıp varsayılmadan)."])
    ws.append(["Statik hesap: hasarların aynı kalacağı varsayılır; hasar eğilimi, enflasyon, IBNR, masraflar ve yeni iş dahil değildir."])
    ws = wb.create_sheet("Segment Etkisi")
    _baslik(ws, ["Senaryo", "Alan", "Değer", "Poliçe", "Mevcut Prim", "Beklenen Prim", "Değişim %", "Mevcut H/P %", "Yeni H/P %"], (26, 18, 22, 8, 18, 18, 10, 11, 11))
    for x in r["segment"]:
        _satir(ws, [x["senaryo"], x["alan"], x["deger"], x["police"], x["prim"], x["yeni"], x["degisim"], x["hp"], x["yeni_hp"]], para_s=(5, 6, 7, 8, 9))
    ws = wb.create_sheet("Poliçe Etkisi")
    _baslik(ws, ["Poliçe No", *segmentler, "Mevcut Prim", "Hasar", *[f"{s['ad']} · Yeni Prim" for s in r["senaryolar"]], *[f"{s['ad']} · Değişim %" for s in r["senaryolar"]]],
            (14, *[16] * len(segmentler), 16, 16, *[20] * (2 * len(r["senaryolar"]))))
    n = len(segmentler)
    for i, p in enumerate(r["policeler"]):
        yeni = [s["detay"][i]["yeni"] for s in r["senaryolar"]]
        _satir(ws, [p["no"], *[p["seg"][a] for a in segmentler], p["prim"], p["hasar"], *yeni, *[degisim(y, p["prim"]) for y in yeni]], para_s=range(n + 2, n + 4 + 2 * len(yeni)))
    uyari_sayfasi(wb, sirala(r["uyarilar"]))
    kaydet(wb, cikti)


def calistir(portfoy_yol: Path, senaryo_yol: Path, cikti: Path, azami_artis: Decimal | None = None, hedef_hp: Decimal | None = None) -> dict:
    policeler, segmentler, uy1 = portfoy_oku(portfoy_yol)
    if not policeler:
        raise ValueError(f"{portfoy_yol.name}: simüle edilecek poliçe yok.")
    senaryolar, uy2 = senaryo_oku(senaryo_yol, segmentler)
    r = simule(policeler, senaryolar, segmentler, azami_artis, hedef_hp)
    r["uyarilar"] = sirala(uy1 + uy2 + r["uyarilar"])
    rapor_yaz(cikti, r, segmentler)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Tarife değişikliği senaryolarının portföy primine ve hasar/prim oranına etkisini simüle eder.")
    p.add_argument("--portfoy", type=Path, default=ORNEK / "portfoy.csv", help="Poliçe No, Prim, Hasar + segment sütunları (Branş, Risk Sınıfı, Bölge ...)")
    p.add_argument("--senaryolar", type=Path, default=ORNEK / "senaryolar.csv", help="Senaryo, Alan, Değer, Değişim %%, Kayıp %%, Açıklama")
    p.add_argument("--azami-artis", help="Poliçe başına azami prim artışı %% (varsayılan: sınır yok)")
    p.add_argument("--hedef-hp", help="Hedef hasar/prim oranı %%; verilirse gerekli genel değişim hesaplanır")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "tarife_etki_simulasyonu.xlsx")
    a = p.parse_args(argv)
    for y in (a.portfoy, a.senaryolar):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    tavan, hedef = (para(x) if x else None for x in (a.azami_artis, a.hedef_hp))
    if (a.azami_artis and (tavan is None or tavan < 0)) or (a.hedef_hp and (hedef is None or hedef <= 0)):
        print("[X] --azami-artis sıfır veya pozitif, --hedef-hp pozitif bir yüzde olmalı")
        return 2
    try:
        r = calistir(a.portfoy, a.senaryolar, a.cikti, tavan, hedef)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    m = r["mevcut"]
    print(f"[OK] Mevcut: {m['police']} poliçe · prim {m['prim']:,.2f} · H/P %{m['hp']}")
    for s in r["senaryolar"]:
        print(f"[OK] {s['ad']}: beklenen prim {s['prim']:,.2f} (%{s['prim_degisim']}) · H/P %{s['hp']}")
    if r["gerekli"] is not None:
        print(f"[OK] Hedef H/P %{hedef} için gerekli genel değişim %{r['gerekli']}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['kim']} · {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
