"""
Dava Risk ve Karşılık Raporu — Workers / Workless kod bloğu
Hukuk › Hukuk Müdürü

Devam eden davaları kaybetme olasılığına göre sınıflar; karşılık, koşullu yükümlülük ve dönem değişimi tablosunu hazırlar.
  Sınıf (şirket davalı)  = olasılık > --muhtemel (%50) Muhtemel · ≤ --uzak (%10) Uzak · arası Mümkün
                           Olasılık sözle de yazılabilir: Muhtemel / Yüksek · Mümkün / Orta · Uzak / Düşük
  Muhtemel               → karşılık ayrılır;  tutar = Tahmini Yükümlülük (boşsa Dava Değeri kullanılır ve uyarılır)
  Mümkün                 → karşılık ayrılmaz; koşullu yükümlülük olarak dipnotta açıklanır
  Uzak                   → karşılık ayrılmaz, açıklama gerekmez
  Şirket davacı          → koşullu varlık; karşılık ayrılmaz, ayrı listelenir
  Dönem değişimi         = karşılık − Önceki Karşılık   (artı: ek karşılık, eksi: konusu kalmayan karşılık)
  Beklenen değer         = tutar × olasılık % (yalnız sayıyla verilen olasılıklarda; yönetim bilgisi içindir, karşılık tutarı değildir)
Sınıflama TMS 37 ve BOBİ FRS'deki genel yaklaşımı izler. Olasılık ve tutar tahmini avukatın, muhasebe kararı mali işlerin ve denetçinindir.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
SUTUN = {"no": ("dosya no", "esas no", "dosya", "no"), "karsi": ("karsi taraf", "davaci davali", "taraf"), "konu": ("konu", "dava konusu", "dava turu", "tur"), "sifat": ("taraf sifati", "sifat", "sirketin sifati", "konum"),
         "deger": ("dava degeri", "talep tutari", "dava tutari", "tutar"), "olasilik": ("kaybetme olasiligi", "olasilik", "aleyhe sonuclanma olasiligi", "risk"), "tahmin": ("tahmini yukumluluk", "tahmini tutar", "avukat tahmini"),
         "onceki": ("onceki karsilik", "mevcut karsilik", "ayrilmis karsilik"), "asama": ("asama", "durum"), "avukat": ("avukat", "sorumlu avukat", "vekil")}
SOZ = {"muhtemel": "Muhtemel", "yuksek": "Muhtemel", "kuvvetle muhtemel": "Muhtemel", "mumkun": "Mümkün", "orta": "Mümkün", "olasi": "Mümkün", "uzak": "Uzak", "dusuk": "Uzak", "zayif": "Uzak"}
ISLEM = {"Muhtemel": "Karşılık ayrılır", "Mümkün": "Koşullu yükümlülük; dipnotta açıklanır", "Uzak": "Karşılık ayrılmaz; açıklama gerekmez", "Değerlendirilmemiş": "Avukat değerlendirmesi bekleniyor",
         "Koşullu varlık": "Karşılık ayrılmaz; lehe dava ayrı izlenir"}
S_RENK = {"Muhtemel": KIRMIZI, "Mümkün": SARI, "Uzak": YESIL, "Değerlendirilmemiş": MAVI}
Z, YUZ = Decimal(0), Decimal(100)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("no", "deger")):
        no = metin(r["no"])
        if not no:
            continue
        deger, o_metin = para(r["deger"]), metin(r["olasilik"])
        if deger is None or deger < 0:
            uy.append(uyari("Yüksek", "Dava okunamadı", f"{no} (satır {r['_satir']})", f"Dava değeri '{metin(r['deger'])}' okunamadı. Değeri belirsiz davalarda 0 yazıp Tahmini Yükümlülük sütununu doldurun. Dava rapora alınmadı."))
            continue
        soz = SOZ.get(katla(o_metin))
        yuzde = None if soz or not o_metin else para(o_metin)
        if o_metin and soz is None and (yuzde is None or not 0 <= yuzde <= 100):
            uy.append(uyari("Orta", "Olasılık okunamadı", no, f"'{o_metin}' okunamadı; 0–100 arası yüzde veya Muhtemel / Mümkün / Uzak yazın. Dava 'Değerlendirilmemiş' sayıldı."))
            yuzde = None
        tahmin, onceki = para(r["tahmin"]), para(r["onceki"])
        l.append({"no": no, "karsi": metin(r["karsi"]), "konu": metin(r["konu"]) or "(konu yok)", "davaci": katla(r["sifat"]).startswith("davaci"), "deger": deger, "soz": soz, "yuzde": yuzde,
                  "tahmin": tahmin if tahmin is not None and tahmin >= 0 else None, "onceki": onceki if onceki and onceki > 0 else Z, "asama": metin(r["asama"]), "avukat": metin(r["avukat"]) or "(avukat yok)"})
    return l, uy


def _kirilim(davalar: list[dict], alan: str) -> list[dict]:
    o = {}
    for d in davalar:
        x = o.setdefault(katla(d[alan]), {"ad": d[alan], "dosya": 0, "deger": Z, "karsilik": Z, "kosullu": Z})
        x["dosya"] += 1
        if not d["davaci"]:
            x["deger"] += d["deger"]
            x["karsilik"] += d["karsilik"]
            x["kosullu"] += d["tutar"] if d["sinif"] == "Mümkün" else Z
    return sorted(o.values(), key=lambda x: (-x["karsilik"], -x["kosullu"], katla(x["ad"])))


def raporla(davalar: list[dict], muhtemel: Decimal = Decimal(50), uzak: Decimal = Decimal(10)) -> dict:
    if not davalar:
        raise ValueError("Raporlanacak dava yok.")
    uy = []
    for d in davalar:
        d["tutar"] = d["tahmin"] if d["tahmin"] is not None else d["deger"]
        d["karsilik"], d["beklenen"] = Z, None
        if d["davaci"]:
            d["sinif"] = "Koşullu varlık"
        elif d["soz"]:
            d["sinif"] = d["soz"]
        elif d["yuzde"] is None:
            d["sinif"] = "Değerlendirilmemiş"
            uy.append(uyari("Yüksek", "Olasılık değerlendirilmemiş", d["no"], f"{d['deger']:,.2f} değerindeki davada kaybetme olasılığı yazılmamış; karşılık kararı verilemez. Avukattan değerlendirme isteyin."))
        else:
            d["sinif"] = "Muhtemel" if d["yuzde"] > muhtemel else "Uzak" if d["yuzde"] <= uzak else "Mümkün"
            d["beklenen"] = yuvarla(d["tutar"] * d["yuzde"] / YUZ)
            if d["yuzde"] == muhtemel:
                uy.append(uyari("Bilgi", "Olasılık eşikte", d["no"], f"Olasılık tam %{muhtemel}; dava 'Mümkün' sınıflandı ve karşılık ayrılmadı. Sınıflamayı avukatla teyit edin."))
        d["islem"] = ISLEM[d["sinif"]]
        if d["sinif"] == "Muhtemel":
            d["karsilik"] = d["tutar"]
            if d["tahmin"] is None:
                uy.append(uyari("Orta", "Tahmini yükümlülük yok", d["no"], f"Kayıp muhtemel ama tahmini yükümlülük yazılmamış; karşılık dava değeri ({d['deger']:,.2f}) üzerinden hesaplandı. Faiz ve masraflar dahil en iyi tahmini avukattan alın."))
            elif d["tahmin"] > d["deger"] > 0:
                uy.append(uyari("Bilgi", "Karşılık dava değerini aşıyor", d["no"], f"Tahmini yükümlülük {d['tahmin']:,.2f}, dava değeri {d['deger']:,.2f}. Faiz, yargılama gideri ve vekâlet ücreti dahilse beklenen bir durumdur."))
        d["degisim"] = Z if d["davaci"] else d["karsilik"] - d["onceki"]
        if not d["davaci"] and d["sinif"] != "Muhtemel" and d["onceki"]:
            uy.append(uyari("Orta", "Önceki karşılığın konusu kalmamış olabilir", d["no"], f"Dava '{d['sinif']}' sınıfında ama önceki dönemden {d['onceki']:,.2f} karşılık var. Sınıflama doğruysa karşılığın iptalini değerlendirin."))
    aleyhe = [d for d in davalar if not d["davaci"]]
    t = lambda sinif: sum((d["tutar"] for d in aleyhe if d["sinif"] == sinif), Z)   # noqa: E731
    siniflar = {s: {"dosya": sum(d["sinif"] == s for d in davalar), "tutar": sum((d["tutar"] for d in davalar if d["sinif"] == s), Z)} for s in ("Muhtemel", "Mümkün", "Uzak", "Değerlendirilmemiş", "Koşullu varlık")}
    toplam = {"deger": sum((d["deger"] for d in aleyhe), Z), "karsilik": t("Muhtemel"), "onceki": sum((d["onceki"] for d in aleyhe), Z), "ek": sum((d["degisim"] for d in aleyhe if d["degisim"] > 0), Z),
              "iptal": sum((-d["degisim"] for d in aleyhe if d["degisim"] < 0), Z), "kosullu": t("Mümkün"), "beklenen": sum((d["beklenen"] for d in aleyhe if d["beklenen"] is not None), Z)}
    sira = {"Muhtemel": 0, "Değerlendirilmemiş": 1, "Mümkün": 2, "Uzak": 3, "Koşullu varlık": 4}
    davalar.sort(key=lambda d: (sira[d["sinif"]], -d["tutar"], d["no"]))
    return {"davalar": davalar, "siniflar": siniflar, "toplam": toplam, "konular": _kirilim(davalar, "konu"), "avukatlar": _kirilim(davalar, "avukat"), "esik": (muhtemel, uzak), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Dava Listesi"
    _baslik(ws, ["Sınıf", "Dosya No", "Karşı Taraf", "Konu", "Şirket", "Aşama", "Avukat", "Dava Değeri", "Olasılık", "Tahmini Yükümlülük", "Karşılık", "Önceki Karşılık", "Dönem Değişimi", "Beklenen Değer", "Yapılacak İşlem"],
            (18, 12, 24, 22, 9, 16, 18, 16, 10, 16, 16, 16, 16, 16, 40))
    for d in r["davalar"]:
        _satir(ws, [d["sinif"], d["no"], d["karsi"], d["konu"], "Davacı" if d["davaci"] else "Davalı", d["asama"], d["avukat"], d["deger"], d["soz"] or (f"%{d['yuzde']}" if d["yuzde"] is not None else ""), d["tahmin"],
                    d["karsilik"] or None, d["onceki"] or None, d["degisim"] or None, d["beklenen"], d["islem"]], para_s=(8, 10, 11, 12, 13, 14), dolgu={1: S_RENK.get(d["sinif"]), 13: SARI if d["degisim"] else None}, sar=(15,))
    ws = wb.create_sheet("Özet")
    _baslik(ws, ["Gösterge", "Dosya", "Tutar"], (52, 8, 18))
    s, t = r["siniflar"], r["toplam"]
    for etiket, dosya, tutar in (("Muhtemel — karşılık ayrılan", s["Muhtemel"]["dosya"], t["karsilik"]), ("Mümkün — koşullu yükümlülük (dipnot)", s["Mümkün"]["dosya"], t["kosullu"]), ("Uzak — açıklama gerekmeyen", s["Uzak"]["dosya"], s["Uzak"]["tutar"]),
                                 ("Değerlendirilmemiş", s["Değerlendirilmemiş"]["dosya"], s["Değerlendirilmemiş"]["tutar"]), ("Şirketin davacı olduğu davalar (koşullu varlık)", s["Koşullu varlık"]["dosya"], s["Koşullu varlık"]["tutar"]),
                                 ("Şirketin davalı olduğu davaların toplam değeri", None, t["deger"]), ("Önceki dönem karşılığı", None, t["onceki"]), ("Ek karşılık (dönem gideri)", None, t["ek"]),
                                 ("Konusu kalmayan karşılık (iptal)", None, t["iptal"]), ("Dönem sonu karşılık", None, t["karsilik"]), ("Beklenen değer toplamı (yönetim bilgisi)", None, t["beklenen"])):
        _satir(ws, [etiket, dosya, tutar], para_s=(3,))
    ws.append([])
    ws.append([f"Sınıf eşikleri: olasılık > %{r['esik'][0]} Muhtemel, ≤ %{r['esik'][1]} Uzak. Tutarlar iskonto edilmemiştir. Karşılık kararı için muhasebe ve denetçi değerlendirmesi gerekir."])
    for ad, veri, baslik in (("Konu Kırılımı", r["konular"], "Konu"), ("Avukat Kırılımı", r["avukatlar"], "Avukat")):
        ws = wb.create_sheet(ad)
        _baslik(ws, [baslik, "Dosya", "Dava Değeri (davalı)", "Karşılık", "Koşullu Yükümlülük"], (30, 8, 20, 18, 20))
        for x in veri:
            _satir(ws, [x["ad"], x["dosya"], x["deger"], x["karsilik"] or None, x["kosullu"] or None], para_s=(3, 4, 5))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, muhtemel: Decimal = Decimal(50), uzak: Decimal = Decimal(10)) -> dict:
    davalar, uy = oku(yol)
    r = raporla(davalar, muhtemel, uzak)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Davaları kaybetme olasılığına göre sınıflayıp karşılık ve koşullu yükümlülük tablosu üretir.")
    p.add_argument("--davalar", type=Path, help="Dosya No, Karşı Taraf, Konu, Taraf Sıfatı, Dava Değeri, Kaybetme Olasılığı, Tahmini Yükümlülük, Önceki Karşılık, Aşama, Avukat")
    p.add_argument("--muhtemel", default="50", help="Olasılık bu yüzdenin üstündeyse Muhtemel (varsayılan 50)")
    p.add_argument("--uzak", default="10", help="Olasılık bu yüzde ve altındaysa Uzak (varsayılan 10)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "dava_risk_karsilik.xlsx")
    a = p.parse_args(argv)
    yol = a.davalar or ORNEK / "davalar.csv"
    muhtemel, uzak = para(a.muhtemel), para(a.uzak)
    if muhtemel is None or uzak is None or not 0 <= uzak < muhtemel < 100:
        print("[X] --uzak sıfır veya pozitif ve --muhtemel değerinden küçük, --muhtemel 100'den küçük olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, muhtemel, uzak)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    s, t = r["siniflar"], r["toplam"]
    print(f"[OK] {len(r['davalar'])} dava · Muhtemel {s['Muhtemel']['dosya']}, Mümkün {s['Mümkün']['dosya']}, Uzak {s['Uzak']['dosya']}, değerlendirilmemiş {s['Değerlendirilmemiş']['dosya']}, lehe {s['Koşullu varlık']['dosya']}")
    print(f"[OK] Karşılık {t['karsilik']:,.2f} (önceki {t['onceki']:,.2f}; ek {t['ek']:,.2f}, iptal {t['iptal']:,.2f}) · koşullu yükümlülük {t['kosullu']:,.2f}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
