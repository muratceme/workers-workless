import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def d(no, deger, yuzde=None, tahmin=None, onceki=0, soz=None, davaci=False, konu="K", avukat="A"):
    n = lambda x: None if x is None else D(str(x))   # noqa: E731
    return {"no": no, "karsi": "", "konu": konu, "davaci": davaci, "deger": D(deger), "soz": soz, "yuzde": n(yuzde), "tahmin": n(tahmin), "onceki": D(onceki), "asama": "", "avukat": avukat}


def tek(*a, **kw):
    secenek = {k: kw.pop(k) for k in list(kw) if k in ("muhtemel", "uzak")}
    r = main.raporla([d("1", *a, **kw)], **secenek)
    x = r["davalar"][0]
    return x["sinif"], x["karsilik"], x["degisim"], x["beklenen"], sorted(u["tur"] for u in r["uyarilar"])


class SinifTesti(unittest.TestCase):
    def test_esikler(self):
        self.assertEqual([tek(1000, y, 800)[:2] for y in (50.01, 50, 10.01, 10, 0, 100)],
                         [("Muhtemel", D(800)), ("Mümkün", D(0)), ("Mümkün", D(0)), ("Uzak", D(0)), ("Uzak", D(0)), ("Muhtemel", D(800))])
        self.assertEqual([tek(1000, soz=s)[:2] for s in ("Muhtemel", "Mümkün", "Uzak")], [("Muhtemel", D(1000)), ("Mümkün", D(0)), ("Uzak", D(0))])
        self.assertEqual((tek(1000, 40, muhtemel=D(30))[0], tek(1000, 8, uzak=D(5))[0]), ("Muhtemel", "Mümkün"))
        self.assertEqual(tek(1000, 50, 800)[4], ["Olasılık eşikte"])

    def test_karsilik_ve_degisim(self):
        self.assertEqual(tek(450000, 70, 380000, 300000), ("Muhtemel", D(380000), D(80000), D("266000.00"), []))
        self.assertEqual(tek(600000, 55), ("Muhtemel", D(600000), D(600000), D("330000.00"), ["Tahmini yükümlülük yok"]))
        self.assertEqual(tek(900000, 10, None, 150000), ("Uzak", D(0), D(-150000), D("90000.00"), ["Önceki karşılığın konusu kalmamış olabilir"]))
        self.assertEqual(tek(1800000, None, 1950000, 1200000, soz="Muhtemel"), ("Muhtemel", D(1950000), D(750000), None, ["Karşılık dava değerini aşıyor"]))
        self.assertEqual(tek(1000, 80, 600, 900), ("Muhtemel", D(600), D(-300), D("480.00"), []))
        self.assertEqual(tek(0, 80, 500), ("Muhtemel", D(500), D(500), D("400.00"), []))
        self.assertEqual(tek(1000, 80, 0), ("Muhtemel", D(0), D(0), D("0.00"), []))
        self.assertEqual(tek(240000), ("Değerlendirilmemiş", D(0), D(0), None, ["Olasılık değerlendirilmemiş"]))
        self.assertEqual(tek(1000, 90, 800, 500, davaci=True), ("Koşullu varlık", D(0), D(0), None, []))

    def test_toplam_ve_kirilim(self):
        l = [d("1", 1000, 60, 800, 500, konu="İş", avukat="Ali"), d("2", 2000, 30, konu="iş", avukat="Ece"), d("3", 500, 5, None, 100, konu="Ticari", avukat="Ali"), d("4", 300, konu="Ticari", avukat="Ece"),
             d("5", 4000, 20, davaci=True, konu="Ticari", avukat="Ali"), d("6", 700, soz="Muhtemel", konu="Vergi", avukat="Ece")]
        r = main.raporla(l)
        self.assertEqual([(x["no"], x["sinif"]) for x in r["davalar"]], [("1", "Muhtemel"), ("6", "Muhtemel"), ("4", "Değerlendirilmemiş"), ("2", "Mümkün"), ("3", "Uzak"), ("5", "Koşullu varlık")])
        self.assertEqual(r["toplam"], {"deger": D(4500), "karsilik": D(1500), "onceki": D(600), "ek": D(1000), "iptal": D(100), "kosullu": D(2000), "beklenen": D("1105.00")})
        self.assertEqual({k: (v["dosya"], v["tutar"]) for k, v in r["siniflar"].items()},
                         {"Muhtemel": (2, D(1500)), "Mümkün": (1, D(2000)), "Uzak": (1, D(500)), "Değerlendirilmemiş": (1, D(300)), "Koşullu varlık": (1, D(4000))})
        self.assertEqual([(x["ad"], x["dosya"], x["deger"], x["karsilik"], x["kosullu"]) for x in r["konular"]], [("İş", 2, D(3000), D(800), D(2000)), ("Vergi", 1, D(700), D(700), D(0)), ("Ticari", 3, D(800), D(0), D(0))])
        self.assertEqual([(x["ad"], x["dosya"], x["deger"], x["karsilik"], x["kosullu"]) for x in r["avukatlar"]], [("Ali", 3, D(1500), D(800), D(0)), ("Ece", 3, D(3000), D(700), D(2000))])
        with self.assertRaises(ValueError):
            main.raporla([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "d.xlsx"
            r = main.calistir(ORNEK / "davalar.csv", cikti)
            self.assertEqual([(x["no"], x["sinif"], x["karsilik"], x["degisim"], x["beklenen"]) for x in r["davalar"]],
                             [("2022/501", "Muhtemel", D(1950000), D(750000), None), ("2025/042", "Muhtemel", D(600000), D(600000), D(330000)), ("2024/118", "Muhtemel", D(380000), D(80000), D(266000)),
                              ("2025/130", "Değerlendirilmemiş", D(0), D(0), None), ("2023/310", "Mümkün", D(0), D(0), D(350000)), ("2025/077", "Uzak", D(0), D(-150000), D(90000)), ("2024/209", "Koşullu varlık", D(0), D(0), None)])
            self.assertEqual(r["toplam"], {"deger": D(6390000), "karsilik": D(2930000), "onceki": D(1650000), "ek": D(1430000), "iptal": D(150000), "kosullu": D(1000000), "beklenen": D(1036000)})
            self.assertEqual([(x["ad"], x["dosya"], x["karsilik"]) for x in r["avukatlar"]], [("Av. Selim Ak", 3, D(1950000)), ("Av. Deniz Er", 4, D(980000))])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Yüksek", "Olasılık değerlendirilmemiş", "2025/130"), ("Orta", "Tahmini yükümlülük yok", "2025/042"), ("Orta", "Önceki karşılığın konusu kalmamış olabilir", "2025/077"),
                              ("Bilgi", "Karşılık dava değerini aşıyor", "2022/501")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Dava Listesi", "Özet", "Konu Kırılımı", "Avukat Kırılımı", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "d.csv"
            y.write_text("Esas No;Taraf;Dava Konusu;Şirketin Sıfatı;Talep Tutarı;Olasılık;Avukat Tahmini;Mevcut Karşılık\n2025/1;A;İş;DAVACI;1.250.000,50;%35;;\n2025/2;B;;;100;YÜKSEK;120;-5\n2025/3;C;;Davalı;100;belki;;\n"
                         "2025/4;D;;;100;150;;\n2025/5;E;;;x;10;;\n2025/6;F;;;0;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["konu"], x["davaci"], x["deger"], x["soz"], x["yuzde"], x["tahmin"], x["onceki"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("2025/1", "İş", True, D("1250000.50"), None, D(35), None, D(0)), ("2025/2", "(konu yok)", False, D(100), "Muhtemel", None, D(120), D(0)), ("2025/3", "(konu yok)", False, D(100), None, None, None, D(0)),
                               ("2025/4", "(konu yok)", False, D(100), None, None, None, D(0)), ("2025/6", "(konu yok)", False, D(0), None, None, None, D(0))],
                              [("Olasılık okunamadı", "2025/3"), ("Olasılık okunamadı", "2025/4"), ("Dava okunamadı", "2025/5 (satır 6)")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--davalar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--muhtemel", "50", "--uzak", "5"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + x) for x in (["--muhtemel", "10"], ["--muhtemel", "100"], ["--uzak=-1"], ["--uzak", "x"])], [2] * 4)
            self.assertEqual(main.main(["--davalar", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Dosya No;Konu\nA;1\n", "Dosya No;Dava Değeri\n;5\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
