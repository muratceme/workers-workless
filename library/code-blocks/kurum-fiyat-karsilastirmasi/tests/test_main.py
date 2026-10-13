import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def f(kurum, kod, fiyat, adet=1, bolge="X", grup="G"):
    return {"kurum": kurum, "bolge": bolge, "grup": grup, "kod": kod, "ad": "", "fiyat": D(fiyat), "adet": D(adet)}


def veri():
    return [f("A", "I1", 100), f("B", "I1", 120), f("C", "I1", 200), f("A", "I2", 50), f("B", "I2", 50), f("C", "I2", 40), f("A", "I3", 10), f("B", "I3", 99)]


class KarsilastirmaTesti(unittest.TestCase):
    def test_endeks(self):
        r = main.karsilastir(veri())
        self.assertEqual([(i["kod"], i["kurum"], i["enk"], i["ortanca"], i["eny"], i["yayilim"], i["enk_kurum"], i["eny_kurum"]) for i in r["islemler"]],
                         [("I1", 3, D(100), D(120), D(200), D("83.33"), "A", "C"), ("I2", 3, D(40), D(50), D(50), D("20.00"), "C", "A"), ("I3", 2, D(10), None, D(99), None, "A", "B")])
        self.assertEqual({(x["kurum"], x["kod"]): (x["endeks"], x["fark"]) for x in r["fiyatlar"]},
                         {("A", "I1"): (D("83.33"), D(-20)), ("B", "I1"): (D("100.00"), D(0)), ("C", "I1"): (D("166.67"), D(80)), ("A", "I2"): (D("100.00"), D(0)), ("B", "I2"): (D("100.00"), D(0)), ("C", "I2"): (D("80.00"), D(-10)),
                          ("A", "I3"): (None, None), ("B", "I3"): (None, None)})
        self.assertEqual([(k["kurum"], k["durum"], k["endeks"], k["islem"], k["karsilastirilan"], k["pahali"], k["ucuz"], k["ustu"], k["eny"]["kod"]) for k in r["kurumlar"]],
                         [("C", "Pahalı", D("141.18"), 2, 2, 1, 0, D(80), "I1"), ("B", "Emsal düzeyinde", D("100.00"), 3, 2, 0, 0, D(0), "I1"), ("A", "Emsal düzeyinde", D("88.24"), 3, 2, 0, 0, D(0), "I2")])
        self.assertEqual((sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]), r["ozet"]),
                         ([("Orta", "Kurum emsallerinden pahalı", "C · X · G"), ("Yüksek", "Aşırı fiyat", "C · I1")], {"kurum": 3, "fiyat": 8, "karsilastirilan": 6, "pahali": 1, "ustu": D(80)}))

    def test_agirlik_ve_esikler(self):
        l = [f("A", "I1", 100, 10), f("B", "I1", 120, 10), f("C", "I1", 200, 1), f("A", "I2", 50, 10), f("B", "I2", 50, 10), f("C", "I2", 40, 9)]
        r = main.karsilastir(l)
        self.assertEqual({k["kurum"]: (k["endeks"], k["durum"], k["ustu"]) for k in r["kurumlar"]}, {"A": (D("88.24"), "Emsal düzeyinde", D(0)), "B": (D("100.00"), "Emsal düzeyinde", D(0)), "C": (D("98.25"), "Emsal düzeyinde", D(80))})
        s = lambda l, **kw: {k["kurum"]: k["durum"] for k in main.karsilastir(l, **kw)["kurumlar"]}   # noqa: E731
        l = [f("A", "I", 100), f("B", "I", 100), f("C", "I", 100), f("P120", "I", 120), f("P121", "I", 121), f("U80", "I", 80), f("U79", "I", 79)]
        self.assertEqual(s(l), {"A": "Emsal düzeyinde", "B": "Emsal düzeyinde", "C": "Emsal düzeyinde", "P120": "Emsal düzeyinde", "P121": "Pahalı", "U80": "Emsal düzeyinde", "U79": "Ucuz"})
        self.assertEqual((s(l, esik=D(110))["P120"], s(l, esik=D(110))["U80"], s(l, asgari=8)["A"]), ("Pahalı", "Ucuz", "Karşılaştırılamadı"))
        a = lambda **kw: [u["kim"] for u in main.karsilastir([f("A", "I", 100), f("B", "I", 100), f("C", "I", 150), f("D", "I", 151), f("E", "I", 100)], **kw)["uyarilar"] if u["tur"] == "Aşırı fiyat"]   # noqa: E731
        self.assertEqual((a(), a(asiri=D(140))), (["D · I"], ["C · I", "D · I"]))
        r = main.karsilastir([f("A", "I", 100), f("B", "I", 200), f("C", "I", 300), f("D", "I", 400), f("A", "I", 10, bolge="Y"), f("E", "I", 20, bolge="Y"), f("G", "I", 30, grup="H")], asgari=2)
        self.assertEqual(([(i["bolge"], i["grup"], i["kurum"], i["ortanca"]) for i in r["islemler"]], [(k["kurum"], k["bolge"], k["endeks"]) for k in r["kurumlar"]], [u["tur"] for u in r["uyarilar"] if u["onem"] == "Bilgi"]),
                         ([("X", "G", 4, D("250.00")), ("X", "H", 1, None), ("Y", "G", 2, D("15.00"))],
                          [("D", "X", D("160.00")), ("E", "Y", D("133.33")), ("C", "X", D("120.00")), ("B", "X", D("80.00")), ("A", "Y", D("66.67")), ("A", "X", D("40.00")), ("G", "X", None)], ["Emsal bulunamadı"]))
        with self.assertRaises(ValueError):
            main.karsilastir([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "kurum_fiyatlari.csv", cikti)
            self.assertEqual([(k["kurum"], k["durum"], k["endeks"], k["pahali"], k["ustu"]) for k in r["kurumlar"]],
                             [("Örnek Hastanesi 4", "Pahalı", D("132.36"), 3, D(709000)), ("Deneme Tıp Merkezi 2", "Emsal düzeyinde", D("105.05"), 0, D(38000)), ("Örnek Hastanesi 3", "Emsal düzeyinde", D("103.51"), 0, D(150000)),
                              ("Örnek Hastanesi 2", "Emsal düzeyinde", D("98.68"), 0, D(61000)), ("Deneme Tıp Merkezi 3", "Emsal düzeyinde", D("98.67"), 0, D(20000)), ("Deneme Tıp Merkezi 1", "Emsal düzeyinde", D("97.07"), 0, D(0)),
                              ("Örnek Hastanesi 1", "Emsal düzeyinde", D("85.81"), 0, D(0)), ("Numune Hastanesi", "Karşılaştırılamadı", None, 0, D(0))])
            self.assertEqual([(i["bolge"], i["kod"], i["kurum"], i["ortanca"]) for i in r["islemler"]],
                             [("Ankara", "ISL-001", 3, D(1500)), ("Ankara", "ISL-002", 3, D(220)), ("Ankara", "ISL-003", 3, D(4800)), ("İstanbul", "ISL-001", 4, D("2300.00")), ("İstanbul", "ISL-002", 4, D("310.00")),
                              ("İstanbul", "ISL-003", 4, D("6700.00")), ("İstanbul", "ISL-004", 4, D("75000.00")), ("İzmir", "ISL-001", 1, None)])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Yüksek", "Aşırı fiyat", "Örnek Hastanesi 4 · ISL-002"), ("Orta", "Kurum emsallerinden pahalı", "Örnek Hastanesi 4 · İstanbul · A Grubu"), ("Bilgi", "Emsal bulunamadı", "Numune Hastanesi · İzmir · A Grubu")])
            self.assertEqual((r["ozet"]["kurum"], r["ozet"]["fiyat"], r["ozet"]["karsilastirilan"], r["ozet"]["pahali"]), (8, 26, 25, 1))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Kurumlar", "İşlemler", "Fiyatlar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "f.csv"
            y.write_text("Hastane;İl;Sınıf;Kod;İşlem;Anlaşma Fiyatı;Yıllık Adet\nA;İst;G1;K1;Muayene;1.250,50;1,5\nB;;;K1;;500;\na;İst;G1;k1;;600;\nC;;;K1;;0;\nD;;;K1;;10;x\n;;;K1;;10;\nE;;;;;10;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["kurum"], x["bolge"], x["grup"], x["kod"], x["ad"], x["fiyat"], x["adet"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("A", "İst", "G1", "K1", "Muayene", D("1250.50"), D("1.5")), ("B", "(bölge yok)", "(grup yok)", "K1", "", D(500), D(1))], [("Fiyat okunamadı", "a · k1"), ("Fiyat okunamadı", "C · K1"), ("Fiyat okunamadı", "D · K1")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--fiyatlar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--esik", "115", "--asiri", "140", "--min", "2"]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--esik", "100"], ["--esik", "200"], ["--asiri", "110"], ["--min", "1"], ["--esik", "x"])], [2] * 5)
            self.assertEqual(main.main(["--fiyatlar", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Kurum;İşlem Kodu\nA;B\n", "Kurum;İşlem Kodu;Fiyat\n;K;5\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
