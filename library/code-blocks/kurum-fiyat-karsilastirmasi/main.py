"""
Kurum Fiyat Karşılaştırması — Workers / Workless kod bloğu
Sağlık Sigortaları › Anlaşmalı Kurumlar Uzmanı

Anlaşmalı sağlık kurumlarının işlem fiyatlarını aynı bölge ve kurum grubundaki emsalleriyle karşılaştırır.
  Emsal kümesi       = aynı bölge + aynı kurum grubu + aynı işlem;   en az --min (3) kurum varsa karşılaştırılır
  Ortanca fiyat      = emsal kümesindeki fiyatların ortancası (çift sayıda ise ortadaki iki fiyatın ortalaması)
  İşlem endeksi      = kurumun fiyatı ÷ ortanca × 100
  Kurum endeksi      = Σ (fiyat × adet) ÷ Σ (ortanca × adet) × 100     (yalnız karşılaştırılabilen işlemler; adet yoksa 1)
  Fiyat yayılımı %   = (en yüksek − en düşük) ÷ ortanca × 100
  Ortanca üstü tutar = Σ max(fiyat − ortanca, 0) × adet
  Durum              = Pahalı (kurum endeksi > --esik, 120) · Ucuz (< 200 − eşik) · Emsal düzeyinde;   aşırı fiyat = işlem endeksi > --asiri (150)
Fiyatların aynı kapsamı (paket, malzeme, hekim ücreti) ifade ettiğini kod bilemez; emsal tanımı sizin verinize göredir.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"kurum": ("kurum", "kurum adi", "saglik kurumu", "hastane"), "bolge": ("bolge", "il", "sehir"), "grup": ("kurum grubu", "grup", "kurum sinifi", "sinif"), "kod": ("islem kodu", "kod", "hizmet kodu"),
         "ad": ("islem adi", "islem", "hizmet adi"), "fiyat": ("fiyat", "anlasma fiyati", "birim fiyat", "tutar"), "adet": ("adet", "islem adedi", "yillik adet")}
DURUM_RENK = {"Pahalı": KIRMIZI, "Ucuz": MAVI, "Emsal düzeyinde": YESIL, "Karşılaştırılamadı": None}


def olcu(x) -> Decimal | None:
    s = metin(x).replace(" ", "")
    if not s:
        return None
    try:
        return Decimal(s.replace(".", "").replace(",", ".") if "," in s else s)
    except InvalidOperation:
        return None


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("kurum", "kod", "fiyat")):
        kurum, kod = metin(r["kurum"]), metin(r["kod"])
        if not kurum or not kod:
            continue
        fiyat, adet = para(r["fiyat"]), para(r["adet"]) if metin(r["adet"]) else Decimal(1)
        anahtar = (katla(kurum), katla(kod))
        if fiyat is None or fiyat <= 0 or adet is None or adet < 0 or anahtar in gorulen:
            uy.append(uyari("Yüksek", "Fiyat okunamadı", f"{kurum} · {kod}", f"Satır {r['_satir']}: " + ("aynı kurum ve işlem daha önce geçti" if anahtar in gorulen else f"fiyat '{metin(r['fiyat'])}' ya da adet '{metin(r['adet'])}' okunamadı")
                            + ". Satır alınmadı."))
            continue
        gorulen.add(anahtar)
        l.append({"kurum": kurum, "bolge": metin(r["bolge"]) or "(bölge yok)", "grup": metin(r["grup"]) or "(grup yok)", "kod": kod, "ad": metin(r["ad"]), "fiyat": fiyat, "adet": adet})
    return l, uy


def _oran(pay: Decimal, payda: Decimal) -> Decimal | None:
    return yuvarla(pay * 100 / payda) if payda else None


def _medyan(d: list[Decimal]) -> Decimal:
    s, n = sorted(d), len(d)
    return s[n // 2] if n % 2 else yuvarla((s[n // 2 - 1] + s[n // 2]) / 2)


def karsilastir(fiyatlar: list[dict], esik: Decimal = Decimal(120), asiri: Decimal = Decimal(150), asgari: int = 3) -> dict:
    if not fiyatlar:
        raise ValueError("Karşılaştırılacak fiyat yok.")
    uy, kume = [], {}
    for f in fiyatlar:
        kume.setdefault((katla(f["bolge"]), katla(f["grup"]), katla(f["kod"])), []).append(f)
    islemler = []
    for g in kume.values():
        fy = [f["fiyat"] for f in g]
        var = len(g) >= asgari
        ort = _medyan(fy) if var else None
        islemler.append({"bolge": g[0]["bolge"], "grup": g[0]["grup"], "kod": g[0]["kod"], "ad": next((f["ad"] for f in g if f["ad"]), ""), "kurum": len(g), "enk": min(fy), "ortanca": ort, "eny": max(fy),
                         "yayilim": _oran(max(fy) - min(fy), ort) if var else None, "eny_kurum": max(g, key=lambda f: f["fiyat"])["kurum"], "enk_kurum": min(g, key=lambda f: f["fiyat"])["kurum"]})
        for f in g:
            f["emsal"], f["ortanca"], f["endeks"] = len(g), ort, _oran(f["fiyat"], ort) if var else None
            f["fark"] = f["fiyat"] - ort if var else None
            if var and f["endeks"] > asiri:
                uy.append(uyari("Yüksek", "Aşırı fiyat", f"{f['kurum']} · {f['kod']}", f"{f['ad'] or f['kod']}: fiyat {f['fiyat']:,.2f}, emsal ortancası {ort:,.2f} (endeks {f['endeks']}; {f['bolge']} · {f['grup']}, {len(g)} kurum)."))
    islemler.sort(key=lambda i: (katla(i["bolge"]), katla(i["grup"]), katla(i["kod"])))
    grup = {}
    for f in fiyatlar:
        grup.setdefault((katla(f["kurum"]), katla(f["bolge"]), katla(f["grup"])), []).append(f)
    kurumlar = []
    for g in grup.values():
        k = [f for f in g if f["endeks"] is not None]
        pay, payda = sum((f["fiyat"] * f["adet"] for f in k), Decimal(0)), sum((f["ortanca"] * f["adet"] for f in k), Decimal(0))
        e = _oran(pay, payda)
        durum = "Karşılaştırılamadı" if e is None else "Pahalı" if e > esik else "Ucuz" if e < 200 - esik else "Emsal düzeyinde"
        kr = {"kurum": g[0]["kurum"], "bolge": g[0]["bolge"], "grup": g[0]["grup"], "islem": len(g), "karsilastirilan": len(k), "endeks": e, "durum": durum, "pahali": sum(1 for f in k if f["endeks"] > esik),
              "ucuz": sum(1 for f in k if f["endeks"] < 200 - esik), "ustu": sum((max(f["fark"], Decimal(0)) * f["adet"] for f in k), Decimal(0)), "eny": max(k, key=lambda f: f["endeks"]) if k else None}
        kurumlar.append(kr)
        kim = f"{kr['kurum']} · {kr['bolge']} · {kr['grup']}"
        if durum == "Pahalı":
            uy.append(uyari("Orta", "Kurum emsallerinden pahalı", kim, f"Kurum endeksi {e} (eşik {esik}); {len(k)} işlemin {kr['pahali']} tanesi eşiğin üstünde. Ortanca üstü tutar {kr['ustu']:,.2f}. "
                            f"En yüksek: {kr['eny']['ad'] or kr['eny']['kod']} (endeks {kr['eny']['endeks']})."))
        elif durum == "Karşılaştırılamadı":
            uy.append(uyari("Bilgi", "Emsal bulunamadı", kim, f"{len(g)} işlemin hiçbirinde aynı bölge ve grupta en az {asgari} kurum yok."))
    kurumlar.sort(key=lambda k: (k["endeks"] is None, -(k["endeks"] or 0), katla(k["kurum"])))
    return {"fiyatlar": sorted(fiyatlar, key=lambda f: (katla(f["bolge"]), katla(f["grup"]), katla(f["kod"]), -f["fiyat"], katla(f["kurum"]))), "islemler": islemler, "kurumlar": kurumlar,
            "ozet": {"kurum": len(kurumlar), "fiyat": len(fiyatlar), "karsilastirilan": sum(1 for f in fiyatlar if f["endeks"] is not None), "pahali": sum(1 for k in kurumlar if k["durum"] == "Pahalı"),
                     "ustu": sum((k["ustu"] for k in kurumlar), Decimal(0))}, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, esik: Decimal, asiri: Decimal) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kurumlar"
    _baslik(ws, ["Kurum", "Bölge", "Kurum Grubu", "Durum", "Kurum Endeksi", "İşlem", "Karşılaştırılan", f"Endeksi > {esik} Olan İşlem", f"Endeksi < {200 - esik} Olan İşlem", "Ortanca Üstü Tutar", "En Yüksek Endeksli İşlem", "Endeksi"],
            (30, 14, 14, 18, 10, 8, 12, 12, 12, 17, 32, 9))
    for k in r["kurumlar"]:
        _satir(ws, [k["kurum"], k["bolge"], k["grup"], k["durum"], k["endeks"], k["islem"], k["karsilastirilan"], k["pahali"] or None, k["ucuz"] or None, k["ustu"] or None, (k["eny"]["ad"] or k["eny"]["kod"]) if k["eny"] else None,
                    k["eny"]["endeks"] if k["eny"] else None], para_s=(5, 10, 12), dolgu={4: DURUM_RENK[k["durum"]]})
    ws.append([])
    ws.append(["Kurum endeksi 100 = emsal ortancası. Adet yazılmışsa işlemler adetle ağırlıklandırılır."])
    ws = wb.create_sheet("İşlemler")
    _baslik(ws, ["Bölge", "Kurum Grubu", "İşlem Kodu", "İşlem Adı", "Kurum Sayısı", "En Düşük", "Ortanca", "En Yüksek", "Yayılım %", "En Düşük Fiyatlı Kurum", "En Yüksek Fiyatlı Kurum"], (14, 14, 12, 32, 9, 14, 14, 14, 10, 28, 28))
    for i in r["islemler"]:
        _satir(ws, [i["bolge"], i["grup"], i["kod"], i["ad"], i["kurum"], i["enk"], i["ortanca"], i["eny"], i["yayilim"], i["enk_kurum"], i["eny_kurum"]], para_s=(6, 7, 8, 9))
    ws = wb.create_sheet("Fiyatlar")
    _baslik(ws, ["Bölge", "Kurum Grubu", "İşlem Kodu", "İşlem Adı", "Kurum", "Fiyat", "Emsal Kurum Sayısı", "Ortanca", "Fark", "İşlem Endeksi", "Adet"], (14, 14, 12, 32, 30, 14, 10, 14, 14, 10, 9))
    for f in r["fiyatlar"]:
        _satir(ws, [f["bolge"], f["grup"], f["kod"], f["ad"], f["kurum"], f["fiyat"], f["emsal"], f["ortanca"], f["fark"], f["endeks"], f["adet"]], para_s=(6, 8, 9, 10),
               dolgu={10: None if f["endeks"] is None else KIRMIZI if f["endeks"] > asiri else SARI if f["endeks"] > esik else MAVI if f["endeks"] < 200 - esik else None})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, esik: Decimal = Decimal(120), asiri: Decimal = Decimal(150), asgari: int = 3) -> dict:
    fiyatlar, uy = oku(yol)
    r = karsilastir(fiyatlar, esik, asiri, asgari)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, esik, asiri)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Anlaşmalı sağlık kurumlarının işlem fiyatlarını bölge ve kurum grubu bazında emsalleriyle karşılaştırır.")
    p.add_argument("--fiyatlar", type=Path, help="Kurum, Bölge, Kurum Grubu, İşlem Kodu, İşlem Adı, Fiyat, Adet")
    p.add_argument("--esik", default="120", help="Pahalı sayılacak endeks; 100 = emsal ortancası (varsayılan 120)")
    p.add_argument("--asiri", default="150", help="Tek işlemde aşırı fiyat sayılacak endeks (varsayılan 150)")
    p.add_argument("--min", type=int, default=3, dest="asgari", help="Karşılaştırma için emsal kümesindeki en az kurum sayısı (varsayılan 3)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kurum_fiyat_karsilastirmasi.xlsx")
    a = p.parse_args(argv)
    yol = a.fiyatlar or ORNEK / "kurum_fiyatlari.csv"
    esik, asiri = olcu(a.esik), olcu(a.asiri)
    if esik is None or not 100 < esik < 200 or asiri is None or asiri < esik or a.asgari < 2:
        print("[X] --esik 100 ile 200 arasında, --asiri eşikten küçük olmamalı, --min en az 2 olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, esik, asiri, a.asgari)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {o['kurum']} kurum · {o['fiyat']} fiyat (karşılaştırılan {o['karsilastirilan']}) · pahalı kurum {o['pahali']} · ortanca üstü tutar {o['ustu']:,.2f}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
