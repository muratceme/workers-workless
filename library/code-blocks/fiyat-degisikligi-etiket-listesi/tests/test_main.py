import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 14)
G = date(2026, 10, 15)


def d(kod, eski, yeni, bozulan=False, tarih=G):
    return {"kod": kod, "ad": kod, "eski": None if eski is None else D(eski), "yeni": D(yeni), "tarih": tarih, "bozulan": bozulan}


class HazirlaTesti(unittest.TestCase):
    def test_indirim_kurali(self):
        gecmis = {"a": [(date(2026, 9, 20), D(90)), (date(2026, 10, 1), D(100))], "b": [(date(2026, 9, 10), D(80))], "c": [(date(2026, 10, 1), D(95))],
                  "t": [(date(2026, 10, 10), D(30)), (date(2026, 10, 13), D(40))]}
        l = [d("A", 100, 85), d("B", 100, 90), d("C", 100, 95), d("T", 40, 35, True), d("Z", 100, 120), d("S", 100, 100), d("Y", None, 10)]
        r = main.hazirla(l, {}, gecmis, BUGUN)
        x = {e["kod"]: e for e in r["etiketler"]}
        # A: son 30 gün (15.09–14.10) en düşük 90 → önceki 90, 85 < 90 gösterilebilir; B: 10.09 pencere dışı → önceki 100
        self.assertEqual([(k, x[k]["onceki"], x[k]["indirim_gosterilir"]) for k in ("A", "B", "C", "T")],
                         [("A", D(90), True), ("B", D(100), True), ("C", D(95), False), ("T", D(40), True)])
        self.assertEqual((x["Z"]["yon"], x["Z"]["degisim"], x["Y"]["yon"], "S" in x), ("Zam", D("20.00"), "Fiyat yok", False))
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"] if u["tur"] != "Raf konumu yok"),
                         [("Bilgi", "Fiyat değişmemiş", "S"), ("Bilgi", "Önceki fiyat eski fiyattan farklı", "A"), ("Yüksek", "İndirim olarak gösterilemez", "C")])
        self.assertEqual(main.hazirla([d("B", 100, 90)], {}, gecmis, BUGUN, gun=40)["etiketler"][0]["onceki"], D(80))

    def test_siralama(self):
        yer = {"a": [{"reyon": "Gıda", "raf": "10", "sira": "1"}], "b": [{"reyon": "Gıda", "raf": "2", "sira": "3"}, {"reyon": "Kasa", "raf": "1", "sira": "1"}],
               "c": [{"reyon": "Gıda", "raf": "2", "sira": "1"}]}
        r = main.hazirla([d("A", 10, 11), d("B", 10, 12), d("C", 10, 13), d("D", 10, 30, tarih=date(2026, 10, 1))], yer, {}, BUGUN)
        self.assertEqual([(e["kod"], e["reyon"], e["raf"]) for e in r["etiketler"]], [("C", "Gıda", "2"), ("B", "Gıda", "2"), ("A", "Gıda", "10"), ("B", "Kasa", "1"), ("D", "", "")])
        self.assertEqual([(z["reyon"], z["etiket"]) for z in r["reyonlar"]], [("Gıda", 3), ("Kasa", 1), ("Konumsuz", 1)])
        self.assertEqual(sorted(u["tur"] for u in r["uyarilar"]), ["Büyük fiyat değişimi", "Geçerlilik tarihi geçmiş", "Raf konumu yok"])
        with self.assertRaises(ValueError):
            main.hazirla([], {}, {}, BUGUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "fiyat_degisiklikleri.csv", cikti, BUGUN, ORNEK / "yerlesim.csv", ORNEK / "fiyat_gecmisi.csv")
            e = r["etiketler"]
            self.assertEqual((len(e), sum(1 for x in e if x["yon"] == "Zam"), sum(1 for x in e if x["yon"] == "İndirim")), (18, 7, 11))
            x = {i["kod"]: i for i in e}
            self.assertEqual((x["8690007"]["onceki"], x["8690011"]["indirim_gosterilir"], x["8690009"]["onceki"], x["8690018"]["onceki"]), (D(159), False, D("39.90"), D(319)))
            self.assertEqual(e[0]["reyon"], "Atıştırmalık")
            self.assertEqual(load_workbook(cikti).sheetnames, ["Etiket Listesi", "Reyon Özeti", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            dy = Path(tmp) / "d.csv"
            dy.write_text("Barkod;Ürün;Mevcut Fiyat;Yeni Fiyat;Tarih\n1;A;10;9;15.10.2026\n1;A;10;9;\n2;B;10;x;\n3;C;0;5;\n4;D;10;5;x\n", encoding="utf-8")
            l, uy = main.degisim_oku(dy)
            self.assertEqual(([x["kod"] for x in l], [u["tur"] for u in uy]), (["1"], ["Ürün iki kez yazılmış", "Satır okunamadı", "Satır okunamadı", "Satır okunamadı"]))
            self.assertEqual(main._siralama("Raf 10") > main._siralama("Raf 2"), True)
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--degisiklikler", str(dy), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "14.10.2026", "--esik", "30", "--gun", "20", "--yerlesim", str(ORNEK / "yerlesim.csv"), "--gecmis", str(ORNEK / "fiyat_gecmisi.csv")]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--esik", "0"], ["--gun", "0"])], [2] * 3)
            self.assertEqual(main.main(arg + ["--yerlesim", str(Path(tmp) / "yok.csv")]), 1)
            dy.write_text("Barkod;Yeni Fiyat\n1;x\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
