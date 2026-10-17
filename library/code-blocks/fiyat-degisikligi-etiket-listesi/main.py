"""
Fiyat Değişikliği Etiket Listesi — Gerisi Bende kod bloğu
Mağaza Operasyonları › Mağaza Müdür Yardımcısı
Merkezden gelen fiyat değişikliklerini reyon, raf ve sıra düzeninde etiket değişim listesine çevirir; indirimli etiketlerde gösterilecek önceki fiyatı hesaplar.
  Değişim %          = (yeni − eski) ÷ eski × 100;  yön: zam / indirim
  Önceki fiyat (indirimde) = geçerlilik gününden önceki 30 gün içinde uygulanan en düşük fiyat (fiyat geçmişi dosyası + eski fiyat)
                       çabuk bozulan ürünlerde (meyve, sebze vb.) indirimden bir önceki fiyat
  İndirim gösterilemez = yeni fiyat, son 30 günün en düşük fiyatından düşük değilse
  Sıralama           = reyon → raf → sıra (yerleşim dosyası); konumu olmayanlar sonda
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
DEGISIM = {"kod": ("urun kodu", "stok kodu", "barkod", "sku"), "ad": ("urun adi", "urun", "aciklama"), "eski": ("eski fiyat", "mevcut fiyat", "onceki fiyat"),
           "yeni": ("yeni fiyat", "fiyat"), "tarih": ("gecerlilik tarihi", "gecerlilik", "baslangic tarihi", "tarih"), "bozulan": ("cabuk bozulan", "taze urun", "bozulabilir")}
YERLESIM = {"kod": ("urun kodu", "stok kodu", "barkod", "sku"), "reyon": ("reyon", "bolum", "kategori"), "raf": ("raf", "raf no", "gondol"), "sira": ("sira", "sira no", "pozisyon")}
GECMIS = {"kod": ("urun kodu", "stok kodu", "barkod", "sku"), "tarih": ("tarih", "gecerlilik tarihi"), "fiyat": ("fiyat", "satis fiyati")}


def degisim_oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, DEGISIM, ("kod", "yeni")):
        kod = metin(r["kod"])
        if not kod:
            continue
        e, y = (para(r[a]) if metin(r[a]) else None for a in ("eski", "yeni"))
        g = tarih(r["tarih"]) if metin(r["tarih"]) else None
        if y is None or y <= 0 or (metin(r["eski"]) and (e is None or e <= 0)) or (metin(r["tarih"]) and g is None):
            uy.append(uyari("Yüksek", "Satır okunamadı", kod, f"Satır {r['_satir']}: eski / yeni fiyat ya da geçerlilik tarihi okunamadı."))
            continue
        if katla(kod) in gorulen:
            uy.append(uyari("Orta", "Ürün iki kez yazılmış", kod, f"Satır {r['_satir']}: ikinci satır alınmadı."))
            continue
        gorulen.add(katla(kod))
        l.append({"kod": kod, "ad": metin(r["ad"]), "eski": e, "yeni": y, "tarih": g, "bozulan": katla(r["bozulan"]) in ("evet", "e", "x", "1")})
    return l, uy


def yerlesim_oku(yol: Path) -> dict:
    d = {}
    for r in kayitlar(yol, YERLESIM, ("kod", "reyon")):
        if metin(r["kod"]):
            d.setdefault(katla(r["kod"]), []).append({"reyon": metin(r["reyon"]), "raf": metin(r["raf"]), "sira": metin(r["sira"])})
    return d


def gecmis_oku(yol: Path) -> dict:
    d = {}
    for r in kayitlar(yol, GECMIS, ("kod", "tarih", "fiyat")):
        g, f = tarih(r["tarih"]), para(r["fiyat"])
        if metin(r["kod"]) and g and f is not None and f > 0:
            d.setdefault(katla(r["kod"]), []).append((g, f))
    return d


def _siralama(x: str):
    return [(0, int(p), "") if p.isdigit() else (1, 0, p) for p in re.findall(r"\d+|\D+", katla(x))] or [(2, 0, "")]


def hazirla(degisimler: list[dict], yerlesim: dict, gecmis: dict, bugun: date, esik: Decimal = Decimal(50), gun: int = 30) -> dict:
    if not degisimler:
        raise ValueError("Fiyat değişikliği satırı yok.")
    uy, etiket = [], []
    for d in degisimler:
        g = d["tarih"] or bugun
        d["degisim"] = yuvarla((d["yeni"] - d["eski"]) * 100 / d["eski"]) if d["eski"] else None
        d["yon"] = "Fiyat yok" if d["eski"] is None else ("İndirim" if d["yeni"] < d["eski"] else ("Zam" if d["yeni"] > d["eski"] else "Değişmedi"))
        d["onceki"], d["indirim_gosterilir"] = None, None
        if d["yon"] == "İndirim":
            gecmisi = sorted(gecmis.get(katla(d["kod"]), []))
            if d["bozulan"]:
                once = [f for t, f in gecmisi if t < g]
                d["onceki"] = once[-1] if once else d["eski"]
            else:
                pencere = [f for t, f in gecmisi if g - timedelta(days=gun) <= t < g] + [d["eski"]]
                d["onceki"] = min(pencere)
            d["indirim_gosterilir"] = d["yeni"] < d["onceki"]
            if not d["indirim_gosterilir"]:
                uy.append(uyari("Yüksek", "İndirim olarak gösterilemez", d["kod"], f"Yeni fiyat {d['yeni']:,.2f}, son {gun} günün en düşük fiyatı {d['onceki']:,.2f}; etikette indirimli / üstü çizili fiyat kullanmayın, normal fiyat etiketi basın."))
            elif d["onceki"] != d["eski"]:
                uy.append(uyari("Bilgi", "Önceki fiyat eski fiyattan farklı", d["kod"], f"Etikette önceki fiyat olarak {d['eski']:,.2f} değil, son {gun} günün en düşük fiyatı {d['onceki']:,.2f} gösterilmeli."))
        if d["yon"] == "Değişmedi":
            uy.append(uyari("Bilgi", "Fiyat değişmemiş", d["kod"], "Eski ve yeni fiyat aynı; etiket değişimi gerekmez."))
        if d["degisim"] is not None and abs(d["degisim"]) >= esik:
            uy.append(uyari("Orta", "Büyük fiyat değişimi", d["kod"], f"%{d['degisim']} değişim ({d['eski']:,.2f} → {d['yeni']:,.2f}); fiyat girişini merkezle doğrulayın."))
        if d["tarih"] and d["tarih"] < bugun:
            uy.append(uyari("Bilgi", "Geçerlilik tarihi geçmiş", d["kod"], f"Fiyat {d['tarih']:%d.%m.%Y} tarihinden beri geçerli; etiket gecikmiş olabilir."))
        konumlar = yerlesim.get(katla(d["kod"])) or [None]
        if konumlar == [None]:
            uy.append(uyari("Orta", "Raf konumu yok", d["kod"], f"{d['ad']}: yerleşim listesinde yok; etiket listesinin sonuna eklendi."))
        if d["yon"] == "Değişmedi":
            continue
        for k in konumlar:
            etiket.append({**d, "reyon": k["reyon"] if k else "", "raf": k["raf"] if k else "", "sira": k["sira"] if k else "", "konum": k is not None})
    etiket.sort(key=lambda e: (not e["konum"], _siralama(e["reyon"]), _siralama(e["raf"]), _siralama(e["sira"]), katla(e["kod"])))
    reyon = {}
    for e in etiket:
        z = reyon.setdefault(e["reyon"] or "Konumsuz", {"reyon": e["reyon"] or "Konumsuz", "etiket": 0, "zam": 0, "indirim": 0})
        z["etiket"] += 1
        z["zam"] += e["yon"] == "Zam"
        z["indirim"] += e["yon"] == "İndirim"
    return {"etiketler": etiket, "reyonlar": list(reyon.values()), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Etiket Listesi"
    _baslik(ws, ["Reyon", "Raf", "Sıra", "Ürün Kodu", "Ürün Adı", "Eski Fiyat", "Yeni Fiyat", "Değişim %", "Yön", "Etikette Önceki Fiyat", "İndirim Gösterilebilir", "Geçerlilik", "Değiştirildi ✓"],
            (14, 7, 6, 14, 30, 11, 11, 9, 9, 12, 10, 11, 10), "D2")
    for e in r["etiketler"]:
        _satir(ws, [e["reyon"], e["raf"], e["sira"], e["kod"], e["ad"], e["eski"], e["yeni"], e["degisim"], e["yon"], e["onceki"],
                    None if e["indirim_gosterilir"] is None else ("Evet" if e["indirim_gosterilir"] else "Hayır"), e["tarih"], ""], para_s=(6, 7, 8, 10), tarih_s=(12,),
               dolgu={9: YESIL if e["yon"] == "İndirim" else (SARI if e["yon"] == "Zam" else None), 11: KIRMIZI if e["indirim_gosterilir"] is False else None, 1: None if e["konum"] else SARI})
    ws = wb.create_sheet("Reyon Özeti")
    _baslik(ws, ["Reyon", "Etiket", "Zam", "İndirim"], (18, 8, 8, 8))
    for z in r["reyonlar"]:
        _satir(ws, [z["reyon"], z["etiket"], z["zam"], z["indirim"]])
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(deg_yol: Path, cikti: Path, bugun: date, yer_yol: Path | None = None, gec_yol: Path | None = None, esik: Decimal = Decimal(50), gun: int = 30) -> dict:
    d, uy = degisim_oku(deg_yol)
    r = hazirla(d, yerlesim_oku(yer_yol) if yer_yol else {}, gecmis_oku(gec_yol) if gec_yol else {}, bugun, esik, gun)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Fiyat değişikliklerini reyon / raf sırasına göre etiket değişim listesine çevirir.")
    p.add_argument("--degisiklikler", type=Path, help="Ürün Kodu, Ürün Adı, Eski Fiyat, Yeni Fiyat, Geçerlilik Tarihi, Çabuk Bozulan")
    p.add_argument("--yerlesim", type=Path, help="Ürün Kodu, Reyon, Raf, Sıra (isteğe bağlı)")
    p.add_argument("--gecmis", type=Path, help="Fiyat geçmişi: Ürün Kodu, Tarih, Fiyat (isteğe bağlı; indirimde son 30 günün en düşük fiyatı için)")
    p.add_argument("--tarih", help="Rapor günü GG.AA.YYYY (varsayılan bugün; örnek veride 14.10.2026)")
    p.add_argument("--esik", default="50", help="Bu yüzdeden büyük değişimler kontrol için işaretlenir (varsayılan 50)")
    p.add_argument("--gun", type=int, default=30, help="İndirimde önceki fiyat için geriye bakılacak gün (varsayılan 30)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "etiket_listesi.xlsx")
    a = p.parse_args(argv)
    ornek = a.degisiklikler is None
    dy = a.degisiklikler or ORNEK / "fiyat_degisiklikleri.csv"
    yy = a.yerlesim or (ORNEK / "yerlesim.csv" if ornek else None)
    gy = a.gecmis or (ORNEK / "fiyat_gecmisi.csv" if ornek else None)
    bugun = tarih(a.tarih) if a.tarih else (date(2026, 10, 14) if ornek else date.today())
    es = para(a.esik)
    if bugun is None or es is None or es <= 0 or a.gun < 1:
        print("[X] --tarih GG.AA.YYYY; --esik 0'dan büyük; --gun en az 1 olmalı")
        return 2
    for y in (dy, yy, gy):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(dy, a.cikti, bugun, yy, gy, es, a.gun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    e = r["etiketler"]
    print(f"[OK] {len(e)} etiket · zam {sum(1 for x in e if x['yon'] == 'Zam')} · indirim {sum(1 for x in e if x['yon'] == 'İndirim')} · indirim olarak gösterilemeyen {sum(1 for x in e if x['indirim_gosterilir'] is False)}")
    for z in r["reyonlar"]:
        print(f"     {z['reyon']}: {z['etiket']} etiket")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Liste: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
