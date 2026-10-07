import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))

import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"


class SayiTarihTesti(unittest.TestCase):
    def test_sayi_bicimleri(self):
        self.assertEqual(main.sayi_coz("1.234,56"), 1234.56)
        self.assertEqual(main.sayi_coz("1,234.56"), 1234.56)
        self.assertEqual(main.sayi_coz("-12.500"), -12500.0)
        self.assertEqual(main.sayi_coz("8,50"), 8.5)
        self.assertEqual(main.sayi_coz("(150,00)"), -150.0)
        self.assertEqual(main.sayi_coz(42), 42.0)

    def test_tarih_bicimleri(self):
        for s in ("2026-09-05", "05.09.2026", "05/09/2026"):
            self.assertEqual(main.tarih_coz(s), date(2026, 9, 5))


class EslestirmeTesti(unittest.TestCase):
    def k(self, kaynak, sira, gun, tutar, aciklama="x"):
        return main.Kayit(kaynak, sira, date(2026, 9, gun), aciklama, tutar)

    def test_referans_tarih_toleransini_asar(self):
        b = [self.k("banka", 1, 10, 500, "Havale 123456789")]
        d = [self.k("defter", 1, 1, 500, "Tahsilat dek 123456789")]
        e, ab, ad = main.eslestir(b, d, gun_toleransi=3)
        self.assertEqual((len(e), e[0].yontem), (1, "Referans"))

    def test_en_yakin_tarih_secilir(self):
        b = [self.k("banka", 1, 10, 100)]
        d = [self.k("defter", 1, 8, 100), self.k("defter", 2, 10, 100)]
        e, _, ad = main.eslestir(b, d)
        self.assertEqual(e[0].defter.sira, 2)
        self.assertEqual([x.sira for x in ad], [1])


class UctanUcaTest(unittest.TestCase):
    def test_ornek_veri(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "m.xlsx"
            e, ab, ad = main.calistir(ORNEK / "banka_ekstresi.csv", ORNEK / "defter_102.csv", cikti)
            self.assertEqual((len(e), len(ab), len(ad)), (9, 6, 4))
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Özet", "Eşleşenler", "Açık - Banka", "Açık - Defter"])
            ozet = {r[0]: r[1] for r in wb["Özet"].iter_rows(min_row=2, values_only=True)}
            self.assertAlmostEqual(ozet["Fark (banka − defter)"], ozet["Açık kalemlerle açıklanan fark"], places=2)
            nedenler = " | ".join(str(r[4]) for r in wb["Açık - Defter"].iter_rows(min_row=2, values_only=True))
            self.assertIn("mükerrer", nedenler)


if __name__ == "__main__":
    unittest.main()
