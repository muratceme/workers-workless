import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 14)


def x(no, asama, kademe="İlkokul", kaynak="Web", ilk=(9, 1), son=None, sonraki=None, kayit=None, kayip_asama=None, neden="", sorumlu="S"):
    t = {k: (date(2026, *v) if v else None) for k, v in (("ilk", ilk), ("son", son), ("sonraki", sonraki), ("kayit", kayit))}
    return {"no": no, "ogrenci": f"Öğrenci {no}", "veli": "", "telefon": "", "kademe": kademe, "sinif": "", "kaynak": kaynak, "asama": asama, "kayip_asama": kayip_asama, **t,
            "son_temas": max(v for k, v in t.items() if k in ("ilk", "son") and v), "neden": neden, "sorumlu": sorumlu}


class AnalizTesti(unittest.TestCase):
    def test_huni_ve_oranlar(self):
        l = [x("1", "Kayıt", kayit=(9, 21)), x("2", "Kayıt", "Lise", kayit=(9, 11)), x("3", "Kayıp", kayip_asama="Ücret görüşmesi", neden="Ücret"), x("4", "Kayıp"),
             x("5", "Okul gezisi", sonraki=(10, 14)), x("6", "Ücret görüşmesi", "Lise", "Fuar", son=(10, 1), sonraki=(10, 5)), x("7", "Ön görüşme", son=(10, 1), sorumlu=""),
             x("8", "Değerlendirme / tanışma", son=(10, 12))]
        r = main.analiz(l, BUGUN)
        self.assertEqual([(h["asama"], h["ulasan"], h["simdi"], h["gecis"]) for h in r["huni"]],
                         [("Ön görüşme", 8, 1, None), ("Okul gezisi", 6, 1, D("75.00")), ("Değerlendirme / tanışma", 5, 1, D("83.33")), ("Ücret görüşmesi", 4, 1, D("80.00")),
                          ("Kayıt", 2, 2, D("50.00"))])
        t = r["toplam"]
        self.assertEqual((t["aday"], t["kayit"], t["kayip"], t["acik"], t["donusum"], t["kapanis"], t["sure"]), (8, 2, 2, 4, D("25.00"), D("50.00"), D("15.0")))
        self.assertEqual([(s["ad"], s["donusum"]) for s in r["kademe"]], [("İlkokul", D("16.67")), ("Lise", D("50.00"))])
        self.assertEqual([(a["no"], a["sebep"]) for a in r["liste"]], [("6", "Takip 9 gün gecikti"), ("5", "Takip günü bugün"), ("7", "Takip tarihi yok, 13 gündür temas yok")])
        self.assertEqual(r["nedenler"], [("Belirtilmemiş", 1), ("Ücret", 1)])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Bilgi", "Kayıp nedeni yok", "4 Öğrenci 4"), ("Bilgi", "Sorumlu yok", "7 Öğrenci 7"), ("Orta", "Takip gecikti", "6 Öğrenci 6")])
        self.assertEqual(len(main.analiz(l, BUGUN, 14)["liste"]), 2)
        with self.assertRaises(ValueError):
            main.analiz([], BUGUN)

    def test_asama(self):
        self.assertEqual([main.asama(s) for s in ("", "Okul ziyareti", "Tanışma sınavı", "Burs görüşmesi", "Kesin kayıt", "Başka okulu seçti", "x")],
                         ["Ön görüşme", "Okul gezisi", "Değerlendirme / tanışma", "Ücret görüşmesi", "Kayıt", "Kayıp", None])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "adaylar.csv", cikti, BUGUN)
            t = r["toplam"]
            self.assertEqual((t["aday"], t["kayit"], t["kayip"], t["acik"], t["donusum"], t["kapanis"], len(r["liste"])), (184, 49, 79, 56, D("26.63"), D("38.28"), 25))
            self.assertEqual(r["huni"][0]["ulasan"], 184)
            self.assertEqual(load_workbook(cikti).sheetnames, ["Takip Listesi", "Huni", "Kademe", "Kaynak", "Sorumlu", "Kayıp Nedenleri", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "a.csv"
            y.write_text("Aday No;Öğrenci;İlk Temas;Durum;Son Aşama;Takip Tarihi\n1;A;01.09.2026;Kayıp;Okul gezisi;\n1;A;;;;\n2;B;x;;;\n3;C;;belki;;\n4;D;;Kayıp;Kayıp;\n5;E;;;;10.10.2026\n",
                         encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(a["no"], a["asama"], a["kayip_asama"], a["sonraki"]) for a in l], [u["tur"] for u in uy]),
                             ([("1", "Kayıp", "Okul gezisi", None), ("5", "Ön görüşme", None, date(2026, 10, 10))],
                              ["Aday iki kez yazılmış", "Satır okunamadı", "Satır okunamadı", "Satır okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            self.assertEqual(main.main(["--adaylar", str(y), "--tarih", "14.10.2026", "--takip", "3", "--cikti", c]), 0)
            self.assertEqual(main.main(["--adaylar", str(y), "--cikti", c]), 0)
            self.assertEqual([main.main(["--cikti", c] + v) for v in (["--tarih", "x"], ["--takip", "0"])], [2, 2])
            self.assertEqual(main.main(["--adaylar", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Öğrenci;Durum\nA;belki\n", encoding="utf-8")
            self.assertEqual(main.main(["--adaylar", str(y), "--cikti", c]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
