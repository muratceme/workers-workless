"""
Aday Öğrenci Takip Listesi — Workers / Workless kod bloğu
Kayıt Kabul ve Öğrenci İşleri › Kayıt Kabul Sorumlusu
Kayıt görüşmesi yapılan aday öğrencileri aşama, kademe ve takip tarihine göre izler; dönüşüm oranlarını ve takip listesini çıkarır.
  Aşamalar      = Ön görüşme → Okul gezisi → Değerlendirme / tanışma → Ücret görüşmesi → Kayıt (Kayıp ayrı)
  Huni          = her aşamaya ulaşan aday sayısı (kayıp adaylarda "Kayıp Aşaması" yazılıysa o aşamaya kadar)
  Dönüşüm %     = kayıt ÷ aday × 100;  kapanış % = kayıt ÷ (kayıt + kayıp) × 100
  Takip listesi = açık adaylardan sonraki takip tarihi bugün ya da geçmiş olanlar, takip tarihi olmayıp son temastan --takip (7) gün geçenler
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
SUTUN = {"no": ("aday no", "kayit no", "no"), "ogrenci": ("ogrenci", "aday", "ogrenci adi", "aday ogrenci"), "veli": ("veli", "veli adi"), "telefon": ("telefon", "veli telefonu", "gsm"),
         "kademe": ("kademe", "okul kademesi", "seviye"), "sinif": ("hedef sinif", "sinif"), "kaynak": ("kaynak", "nereden duydu", "basvuru kaynagi"),
         "ilk": ("ilk temas tarihi", "ilk temas", "basvuru tarihi"), "asama": ("asama", "durum"), "kayip_asama": ("kayip asamasi", "son asama"),
         "son": ("son temas tarihi", "son temas", "son gorusme"), "sonraki": ("sonraki takip tarihi", "sonraki takip", "takip tarihi"), "neden": ("kayip nedeni", "neden"),
         "sorumlu": ("sorumlu", "temsilci", "danisman"), "kayit": ("kayit tarihi",)}
ASAMALAR = ("Ön görüşme", "Okul gezisi", "Değerlendirme / tanışma", "Ücret görüşmesi", "Kayıt")


def asama(x) -> str | None:
    k = katla(x)
    if not k or k in ("on gorusme", "ilk gorusme", "bilgi alma", "telefon", "basvuru"):
        return "Ön görüşme"
    if "gezi" in k or "ziyaret" in k or "tanitim" in k:
        return "Okul gezisi"
    if "degerlendirme" in k or "tanisma" in k or "sinav" in k or "mulakat" in k or "gozlem" in k:
        return "Değerlendirme / tanışma"
    if "ucret" in k or "teklif" in k or "indirim" in k or "burs" in k or "sozlesme" in k:
        return "Ücret görüşmesi"
    if k in ("kayit", "kayit oldu", "kesin kayit", "kayitli", "kaydoldu"):
        return "Kayıt"
    if "kayip" in k or "vazgecti" in k or "iptal" in k or "baska okul" in k or "olumsuz" in k:
        return "Kayıp"
    return None


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("ogrenci", "asama")):
        if not metin(r["ogrenci"]):
            continue
        kim = metin(r["no"]) or metin(r["ogrenci"])
        a = asama(r["asama"])
        ka = asama(r["kayip_asama"]) if metin(r["kayip_asama"]) else None
        t = {k: (tarih(r[k]) if metin(r[k]) else None) for k in ("ilk", "son", "sonraki", "kayit")}
        if a is None or any(metin(r[k]) and v is None for k, v in t.items()) or (metin(r["kayip_asama"]) and ka in (None, "Kayıp")):
            uy.append(uyari("Yüksek", "Satır okunamadı", kim, f"Satır {r['_satir']}: aşama ('{metin(r['asama'])}'), kayıp aşaması ya da tarihlerden biri okunamadı."))
            continue
        if metin(r["no"]) and katla(r["no"]) in gorulen:
            uy.append(uyari("Orta", "Aday iki kez yazılmış", kim, f"Satır {r['_satir']}: ikinci satır alınmadı."))
            continue
        if metin(r["no"]):
            gorulen.add(katla(r["no"]))
        son = max((v for k, v in t.items() if k in ("ilk", "son") and v), default=None)
        l.append({"no": metin(r["no"]), "ogrenci": metin(r["ogrenci"]), "veli": metin(r["veli"]), "telefon": metin(r["telefon"]), "kademe": metin(r["kademe"]) or "Belirtilmemiş",
                  "sinif": metin(r["sinif"]), "kaynak": metin(r["kaynak"]) or "Belirtilmemiş", "asama": a, "kayip_asama": ka, **t, "son_temas": son, "neden": metin(r["neden"]),
                  "sorumlu": metin(r["sorumlu"])})
    return l, uy


def _oran(a: int, b: int) -> Decimal | None:
    return yuvarla(Decimal(a) * 100 / b) if b else None


def _ulasilan(x: dict) -> int:
    if x["asama"] == "Kayıp":
        return ASAMALAR.index(x["kayip_asama"]) if x["kayip_asama"] else 0
    return ASAMALAR.index(x["asama"])


def analiz(adaylar: list[dict], bugun: date, takip: int = 7) -> dict:
    if not adaylar:
        raise ValueError("Okunabilen aday yok.")
    uy, liste = [], []
    for x in adaylar:
        kim = f"{x['no']} {x['ogrenci']}".strip()
        x["acik"] = x["asama"] not in ("Kayıt", "Kayıp")
        x["gun"] = (bugun - x["son_temas"]).days if x["son_temas"] else None
        x["sebep"] = ""
        if x["asama"] == "Kayıp" and not x["neden"]:
            uy.append(uyari("Bilgi", "Kayıp nedeni yok", kim, "Kayıp adayda neden yazılmamış."))
        if not x["acik"]:
            continue
        if not x["sorumlu"]:
            uy.append(uyari("Bilgi", "Sorumlu yok", kim, "Açık adayın takip sorumlusu yazılmamış."))
        if x["sonraki"] and x["sonraki"] <= bugun:
            gecikme = (bugun - x["sonraki"]).days
            x["sebep"] = "Takip günü bugün" if gecikme == 0 else f"Takip {gecikme} gün gecikti"
            if gecikme > 3:
                uy.append(uyari("Orta", "Takip gecikti", kim, f"{x['sorumlu'] or 'Sorumlu yok'}: takip {x['sonraki']:%d.%m.%Y} için planlanmıştı ({x['asama']})."))
        elif not x["sonraki"] and x["gun"] is not None and x["gun"] >= takip:
            x["sebep"] = f"Takip tarihi yok, {x['gun']} gündür temas yok"
        if x["sebep"]:
            liste.append(x)
    liste.sort(key=lambda x: (-ASAMALAR.index(x["asama"]), x["sonraki"] or date.min, katla(x["ogrenci"])))

    def ozet(fn):
        g = {}
        for x in adaylar:
            g.setdefault(fn(x), []).append(x)
        satir = []
        for ad, l in g.items():
            kayit = [x for x in l if x["asama"] == "Kayıt"]
            kayip = sum(1 for x in l if x["asama"] == "Kayıp")
            sure = [(x["kayit"] - x["ilk"]).days for x in kayit if x["kayit"] and x["ilk"] and x["kayit"] >= x["ilk"]]
            satir.append({"ad": ad, "aday": len(l), "kayit": len(kayit), "kayip": kayip, "acik": sum(1 for x in l if x["acik"]), "donusum": _oran(len(kayit), len(l)),
                          "kapanis": _oran(len(kayit), len(kayit) + kayip), "sure": yuvarla(Decimal(sum(sure)) / len(sure), "0.1") if sure else None})
        return sorted(satir, key=lambda s: (-s["aday"], katla(s["ad"])))
    huni = []
    for i, a in enumerate(ASAMALAR):
        n = sum(1 for x in adaylar if _ulasilan(x) >= i)
        huni.append({"asama": a, "ulasan": n, "simdi": sum(1 for x in adaylar if x["asama"] == a), "gecis": _oran(n, huni[-1]["ulasan"]) if huni else None})
    nedenler = {}
    for x in adaylar:
        if x["asama"] == "Kayıp":
            nedenler[x["neden"] or "Belirtilmemiş"] = nedenler.get(x["neden"] or "Belirtilmemiş", 0) + 1
    return {"liste": liste, "kademe": ozet(lambda x: x["kademe"]), "kaynak": ozet(lambda x: x["kaynak"]), "sorumlu": ozet(lambda x: x["sorumlu"] or "Atanmamış"),
            "toplam": ozet(lambda x: "TOPLAM")[0], "huni": huni, "nedenler": sorted(nedenler.items(), key=lambda n: (-n[1], katla(n[0]))), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Takip Listesi"
    _baslik(ws, ["Aday No", "Öğrenci", "Veli", "Telefon", "Kademe", "Hedef Sınıf", "Aşama", "Son Temas", "Sonraki Takip", "Neden Listede", "Sorumlu"], (9, 20, 20, 14, 11, 9, 20, 11, 11, 32, 14), "C2")
    for x in r["liste"]:
        _satir(ws, [x["no"], x["ogrenci"], x["veli"], x["telefon"], x["kademe"], x["sinif"], x["asama"], x["son_temas"], x["sonraki"], x["sebep"], x["sorumlu"]], tarih_s=(8, 9),
               dolgu={10: SARI if "gecikti" in x["sebep"] else None})
    ws = wb.create_sheet("Huni")
    _baslik(ws, ["Aşama", "Ulaşan Aday", "Şu An Bu Aşamada", "Önceki Aşamadan Geçiş %"], (24, 11, 13, 14))
    for h in r["huni"]:
        _satir(ws, [h["asama"], h["ulasan"], h["simdi"], h["gecis"]], para_s=(4,))
    ws.append([])
    ws.append(["Kayıp adaylar, 'Kayıp Aşaması' yazılıysa o aşamaya kadar, yazılı değilse yalnız ön görüşmede sayılır."])
    for ad, l in (("Kademe", r["kademe"]), ("Kaynak", r["kaynak"]), ("Sorumlu", r["sorumlu"])):
        ws = wb.create_sheet(ad)
        _baslik(ws, [ad, "Aday", "Kayıt", "Kayıp", "Açık", "Dönüşüm %", "Kapanış %", "Ort. Kayıt Süresi (gün)"], (22, 7, 7, 7, 7, 10, 10, 12))
        for s in l + [r["toplam"]]:
            n = _satir(ws, [s["ad"], s["aday"], s["kayit"], s["kayip"], s["acik"], s["donusum"], s["kapanis"], s["sure"]], para_s=(6, 7, 8))
            if s["ad"] == "TOPLAM":
                for h in ws[n]:
                    h.font = Font(bold=True)
    ws = wb.create_sheet("Kayıp Nedenleri")
    _baslik(ws, ["Neden", "Aday"], (36, 8))
    for n, a in r["nedenler"]:
        _satir(ws, [n, a])
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, takip: int = 7) -> dict:
    adaylar, uy = oku(yol)
    r = analiz(adaylar, bugun, takip)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Aday öğrencileri aşama ve takip tarihine göre izler, dönüşüm oranlarını hesaplar.")
    p.add_argument("--adaylar", type=Path, help="Aday No, Öğrenci, Veli, Telefon, Kademe, Hedef Sınıf, Kaynak, İlk Temas Tarihi, Aşama, Kayıp Aşaması, Son Temas Tarihi, "
                                                "Sonraki Takip Tarihi, Kayıp Nedeni, Sorumlu, Kayıt Tarihi")
    p.add_argument("--tarih", help="Rapor günü GG.AA.YYYY (varsayılan bugün; örnek veride 14.10.2026)")
    p.add_argument("--takip", type=int, default=7, help="Takip tarihi olmayan açık adayda kaç gün temassızlıkta listeye alınır (varsayılan 7)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "aday_ogrenci_takip.xlsx")
    a = p.parse_args(argv)
    yol = a.adaylar or ORNEK / "adaylar.csv"
    bugun = tarih(a.tarih) if a.tarih else (date(2026, 10, 14) if a.adaylar is None else date.today())
    if bugun is None or a.takip < 1:
        print("[X] --tarih GG.AA.YYYY, --takip en az 1 olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, bugun, a.takip)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {t['aday']} aday · {t['kayit']} kayıt · {t['kayip']} kayıp · {t['acik']} açık · dönüşüm %{t['donusum']} · kapanış %{t['kapanis']} · takip listesinde {len(r['liste'])}")
    for s in r["kademe"]:
        print(f"     {s['ad']}: {s['aday']} aday · {s['kayit']} kayıt (%{s['donusum']})")
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Liste: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
