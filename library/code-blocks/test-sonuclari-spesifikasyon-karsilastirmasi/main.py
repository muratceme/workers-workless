"""
Test Sonuçları Spesifikasyon Karşılaştırması — Workers / Workless kod bloğu
Ar-Ge ve Ürün Geliştirme › Ar-Ge Mühendisi

Prototip test sonuçlarını spesifikasyon limitleriyle karşılaştırır; uygunsuzlukları ve prototipler arası eğilimi raporlar.
  Uygun      = alt limit ≤ sonuç ≤ üst limit                (limitin kendisi uygundur; tek taraflı limit olabilir)
  Marj       = ortalamanın en yakın limite uzaklığı          (ortalama limit dışındaysa eksi)
  Marj %     = marj ÷ (üst − alt) × 100;  tek taraflı limitte marj ÷ |limit| × 100
  Eğilim     = son prototipin marj %'si − ilk prototipin marj %'si;  ±--egilim-esik (1 puan) içindeyse "Değişmedi"
  Dikkat     = eğilim kötüleşiyor ve son marj % < --marj-esik (10), veya son prototipte uygunsuz sonuç var
Prototip sırası dosyadaki ilk görülme sırasıdır. Az sayıda numuneyle eğilim yalnız gösterge niteliğindedir; istatistiksel anlamlılık testi yapılmaz.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SPEK_SUTUN = {"par": ("parametre", "test", "ozellik", "olcut"), "alt": ("alt limit", "min", "en az"), "ust": ("ust limit", "max", "maks", "en cok"), "hedef": ("hedef", "nominal"),
              "birim": ("birim",)}
SONUC_SUTUN = {"proto": ("prototip", "versiyon", "revizyon", "surum"), "numune": ("numune", "numune no", "tekrar"), "par": ("parametre", "test", "ozellik", "olcut"),
               "deger": ("sonuc", "deger", "olculen", "test sonucu"), "tarih": ("tarih",), "kosul": ("kosul", "test kosulu", "not")}
Z = Decimal(0)


def sayi(x) -> Decimal | None:
    """Virgül ve nokta ondalık ayırıcıdır; binlik ayırıcı kabul edilmez."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    try:
        return Decimal(re.sub(r"[^\d,.\-+]", "", str(x).replace("−", "-")).replace(",", "."))
    except InvalidOperation:
        return None


def oku(sonuc_yol: Path, spek_yol: Path) -> tuple[list[dict], dict, list[dict]]:
    uy, spek, l = [], {}, []
    for r in kayitlar(spek_yol, SPEK_SUTUN, ("par",)):
        p = metin(r["par"])
        if not p:
            continue
        alt, ust, hedef = sayi(r["alt"]), sayi(r["ust"]), sayi(r["hedef"])
        if alt is None and ust is None:
            uy.append(uyari("Yüksek", "Limit tanımı eksik", p, f"{spek_yol.name} satır {r['_satir']}: alt veya üst limitten en az biri girilmeli. Parametre karşılaştırılamadı."))
        elif alt is not None and ust is not None and alt >= ust:
            uy.append(uyari("Yüksek", "Alt limit üst limitten küçük değil", p, f"{spek_yol.name} satır {r['_satir']}: alt {alt}, üst {ust}. Parametre karşılaştırılamadı."))
        elif katla(p) in spek:
            uy.append(uyari("Yüksek", "Parametre tekrar ediyor", p, f"{spek_yol.name} satır {r['_satir']}: ilk satır kullanıldı."))
        else:
            spek[katla(p)] = {"par": p, "alt": alt, "ust": ust, "hedef": hedef, "birim": metin(r["birim"])}
    if not spek:
        raise ValueError(f"{spek_yol.name}: kullanılabilir spesifikasyon satırı yok.")
    for r in kayitlar(sonuc_yol, SONUC_SUTUN, ("par", "deger")):
        if not metin(r["par"]):
            continue
        v = sayi(r["deger"])
        if v is None:
            uy.append(uyari("Yüksek", "Sonuç okunamadı", f"{metin(r['proto'])} · {metin(r['par'])}".strip(" ·"), f"{sonuc_yol.name} satır {r['_satir']}: '{metin(r['deger'])}'. Sonuç alınmadı."))
            continue
        l.append({"proto": metin(r["proto"]) or "—", "numune": metin(r["numune"]), "par": metin(r["par"]), "deger": v, "tarih": tarih(r["tarih"]), "kosul": metin(r["kosul"])})
    return l, spek, uy


def uygun_mu(v: Decimal, s: dict) -> str:
    return "Uygunsuz (üst)" if s["ust"] is not None and v > s["ust"] else "Uygunsuz (alt)" if s["alt"] is not None and v < s["alt"] else "Uygun"


def marj(ort: Decimal, s: dict) -> tuple[Decimal, Decimal]:
    if s["alt"] is not None and s["ust"] is not None:
        m = min(ort - s["alt"], s["ust"] - ort)
        return m, yuvarla(m / (s["ust"] - s["alt"]) * 100)
    lim = s["ust"] if s["ust"] is not None else s["alt"]
    m = s["ust"] - ort if s["ust"] is not None else ort - s["alt"]
    return m, yuvarla(m / abs(lim) * 100) if lim else Z


def karsilastir(sonuclar: list[dict], spek: dict, egilim_esik: Decimal = Decimal(1), marj_esik: Decimal = Decimal(10)) -> dict:
    uy, eslesmeyen, l = [], defaultdict(int), []
    for x in sonuclar:
        s = spek.get(katla(x["par"]))
        if not s:
            eslesmeyen[x["par"]] += 1
            continue
        l.append(x | {"s": s, "sonuc": uygun_mu(x["deger"], s)})
    for p, n in eslesmeyen.items():
        uy.append(uyari("Yüksek", "Parametrenin spesifikasyonu yok", p, f"{n} sonuç için spesifikasyon dosyasında karşılık yok; karşılaştırılmadı."))
    if not l:
        raise ValueError("Spesifikasyonu tanımlı parametreye ait sonuç yok.")
    protolar = list(dict.fromkeys(x["proto"] for x in l))
    ozet = []
    for k, s in spek.items():
        for pr in protolar:
            d = [x for x in l if katla(x["par"]) == k and x["proto"] == pr]
            if not d:
                continue
            v = [x["deger"] for x in d]
            ort = sum(v) / len(v)
            m, mp = marj(ort, s)
            ozet.append({"par": s["par"], "anahtar": k, "proto": pr, "s": s, "n": len(v), "ort": yuvarla(ort, "0.0001"), "min": min(v), "max": max(v), "marj": m, "marj_oran": mp,
                         "uygunsuz": sum(x["sonuc"] != "Uygun" for x in d), "hedef_fark": ort - s["hedef"] if s["hedef"] is not None else None})
    egilim = []
    for k, s in spek.items():
        o = [x for x in ozet if x["anahtar"] == k]
        if not o:
            uy.append(uyari("Orta", "Parametre test edilmemiş", s["par"], "Spesifikasyonda var, sonuç dosyasında hiç sonucu yok."))
            continue
        fark = o[-1]["marj_oran"] - o[0]["marj_oran"] if len(o) > 1 else None
        yon = "Tek prototip" if fark is None else "İyileşiyor" if fark > egilim_esik else "Kötüleşiyor" if fark < -egilim_esik else "Değişmedi"
        neden = []
        if o[-1]["uygunsuz"]:
            neden.append(f"son prototipte {o[-1]['uygunsuz']} uygunsuz sonuç")
        if yon == "Kötüleşiyor" and o[-1]["marj_oran"] < marj_esik:
            neden.append(f"marj daralıyor (son marj %{o[-1]['marj_oran']})")
        egilim.append({"par": s["par"], "s": s, "protolar": o, "ilk": o[0]["marj_oran"], "son": o[-1]["marj_oran"], "fark": fark, "yon": yon, "dikkat": "; ".join(neden)})
    proto = []
    for pr in protolar:
        d = [x for x in l if x["proto"] == pr]
        up = sorted({x["par"] for x in d if x["sonuc"] != "Uygun"})
        proto.append({"proto": pr, "sonuc_sayisi": len(d), "parametre": len({katla(x["par"]) for x in d}), "uygunsuz": sum(x["sonuc"] != "Uygun" for x in d), "uygunsuz_par": up,
                      "durum": "Uygunsuzluk var" if up else "Tüm sonuçlar uygun"})
    return {"sonuclar": l, "ozet": ozet, "egilim": egilim, "proto": proto, "protolar": protolar, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    yrenk = {"İyileşiyor": YESIL, "Kötüleşiyor": KIRMIZI, "Değişmedi": None, "Tek prototip": None}
    ws = wb.active
    ws.title = "Prototip Özeti"
    _baslik(ws, ["Prototip", "Parametre Sayısı", "Sonuç Sayısı", "Uygunsuz Sonuç", "Durum", "Uygunsuz Parametreler"], (16, 10, 10, 10, 20, 70))
    for p in r["proto"]:
        _satir(ws, [p["proto"], p["parametre"], p["sonuc_sayisi"], p["uygunsuz"] or None, p["durum"], ", ".join(p["uygunsuz_par"])], dolgu={5: KIRMIZI if p["uygunsuz_par"] else YESIL}, sar=(6,))
    ws = wb.create_sheet("Eğilim")
    _baslik(ws, ["Parametre", "Birim", "Alt Limit", "Üst Limit", "Hedef", *[f"{p} · ortalama" for p in r["protolar"]], "İlk Marj %", "Son Marj %", "Değişim (puan)", "Eğilim", "Dikkat",
                 "Yorum (siz doldurun)"], (26, 8, 10, 10, 10, *[14] * len(r["protolar"]), 10, 10, 10, 13, 44, 30))
    n = len(r["protolar"])
    for e in r["egilim"]:
        ort = {o["proto"]: o["ort"] for o in e["protolar"]}
        _satir(ws, [e["par"], e["s"]["birim"], e["s"]["alt"], e["s"]["ust"], e["s"]["hedef"], *[ort.get(p) for p in r["protolar"]], e["ilk"], e["son"], e["fark"], e["yon"], e["dikkat"], ""],
               para_s=(6 + n, 7 + n, 8 + n), dolgu={9 + n: yrenk[e["yon"]], 10 + n: SARI if e["dikkat"] else None}, sar=(10 + n,))
    ws.append([])
    ws.append(["Marj: ortalamanın en yakın limite uzaklığı (limit dışındaysa eksi). Eğilim, ilk ve son prototipin marj yüzdeleri arasındaki farktır; istatistiksel test değildir."])
    ws = wb.create_sheet("Parametre Özeti")
    _baslik(ws, ["Parametre", "Prototip", "Numune Sayısı", "Ortalama", "En Küçük", "En Büyük", "Alt Limit", "Üst Limit", "Marj", "Marj %", "Hedeften Fark", "Uygunsuz Sonuç"],
            (26, 14, 9, 12, 11, 11, 10, 10, 11, 9, 12, 10))
    for o in r["ozet"]:
        _satir(ws, [o["par"], o["proto"], o["n"], o["ort"], o["min"], o["max"], o["s"]["alt"], o["s"]["ust"], o["marj"], o["marj_oran"], o["hedef_fark"], o["uygunsuz"] or None], para_s=(10,),
               dolgu={12: KIRMIZI if o["uygunsuz"] else None, 10: KIRMIZI if o["marj_oran"] < 0 else None})
    ws = wb.create_sheet("Sonuçlar")
    _baslik(ws, ["Prototip", "Numune", "Parametre", "Sonuç", "Birim", "Alt Limit", "Üst Limit", "Değerlendirme", "Tarih", "Koşul"], (14, 10, 26, 12, 8, 10, 10, 16, 12, 30))
    for x in sorted(r["sonuclar"], key=lambda x: x["sonuc"] == "Uygun"):
        _satir(ws, [x["proto"], x["numune"], x["par"], x["deger"], x["s"]["birim"], x["s"]["alt"], x["s"]["ust"], x["sonuc"], x["tarih"], x["kosul"]], tarih_s=(9,),
               dolgu={8: KIRMIZI if x["sonuc"] != "Uygun" else None})
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(sonuc_yol: Path, spek_yol: Path, cikti: Path, **esik) -> dict:
    sonuclar, spek, uy = oku(sonuc_yol, spek_yol)
    r = karsilastir(sonuclar, spek, **esik)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, r["uyarilar"])
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Prototip test sonuçlarını spesifikasyon limitleriyle karşılaştırır, eğilimi raporlar.")
    p.add_argument("--sonuclar", type=Path, default=ORNEK / "test_sonuclari.csv", help="Prototip, Numune, Parametre, Sonuç, Tarih, Koşul")
    p.add_argument("--spesifikasyon", type=Path, default=ORNEK / "spesifikasyon.csv", help="Parametre, Alt Limit, Üst Limit, Hedef, Birim")
    p.add_argument("--egilim-esik", default="1", help="Bu kadar yüzde puanın altındaki marj değişimi 'Değişmedi' sayılır (varsayılan 1)")
    p.add_argument("--marj-esik", default="10", help="Kötüleşen parametrede bu marj %%'sinin altı 'Dikkat' olarak işaretlenir (varsayılan 10)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "test_spesifikasyon.xlsx")
    a = p.parse_args(argv)
    for y in (a.sonuclar, a.spesifikasyon):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    ee, me = para(a.egilim_esik), para(a.marj_esik)
    if ee is None or me is None or ee < 0 or me < 0:
        print("[X] --egilim-esik ve --marj-esik sıfır veya pozitif sayı olmalı")
        return 2
    try:
        r = calistir(a.sonuclar, a.spesifikasyon, a.cikti, egilim_esik=ee, marj_esik=me)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    for x in r["proto"]:
        print(f"[OK] {x['proto']}: {x['sonuc_sayisi']} sonuç · {x['uygunsuz']} uygunsuz" + (f" ({', '.join(x['uygunsuz_par'])})" if x["uygunsuz_par"] else ""))
    for e in r["egilim"]:
        if e["dikkat"]:
            print(f"[!] {e['par']} · {e['yon']}: {e['dikkat']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
