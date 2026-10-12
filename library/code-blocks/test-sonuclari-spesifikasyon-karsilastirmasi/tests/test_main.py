import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def sp(par="Akım", alt=None, ust=None, hedef=None):
    f = lambda v: None if v is None else D(str(v))   # noqa: E731
    return {"par": par, "alt": f(alt), "ust": f(ust), "hedef": f(hedef), "birim": ""}


def so(proto, deger, par="Akım"):
    return {"proto": proto, "numune": "", "par": par, "deger": D(str(deger)), "tarih": None, "kosul": ""}


SPEK = {"akim": sp(alt=2, ust=3, hedef=2.5), "gurultu": sp("Gürültü", ust=60), "debi": sp("Debi", alt=40)}


class KarsilastirmaTesti(unittest.TestCase):
    def test_uygunluk_ve_marj(self):
        s = SPEK["akim"]
        self.assertEqual([main.uygun_mu(D(str(v)), s) for v in (1.99, 2, 2.5, 3, 3.01)], ["Uygunsuz (alt)", "Uygun", "Uygun", "Uygun", "Uygunsuz (üst)"])
        self.assertEqual([main.marj(D(str(v)), s) for v in (2.75, 3, 3.2)], [(D("0.25"), D("25.00")), (D(0), D("0.00")), (D("-0.2"), D("-20.00"))])
        self.assertEqual(main.marj(D("59.5"), SPEK["gurultu"]), (D("0.5"), D("0.83")))
        self.assertEqual(main.marj(D("38.5"), SPEK["debi"]), (D("-1.5"), D("-3.75")))

    def test_egilim(self):
        l = [so("P1", 2.9), so("P1", 3.1), so("P2", 2.7), so("P2", 2.8), so("P3", 2.5), so("P3", 2.6), so("P1", 55, "Gürültü"), so("P2", 57, "Gürültü"), so("P3", 59.5, "Gürültü"),
             so("P1", 50, "Debi"), so("P3", 50.2, "Debi")]
        r = main.karsilastir(l, SPEK)
        e = {x["par"]: (x["ilk"], x["son"], x["fark"], x["yon"], x["dikkat"]) for x in r["egilim"]}
        self.assertEqual(e["Akım"], (D("0.00"), D("45.00"), D("45.00"), "İyileşiyor", ""))
        self.assertEqual(e["Gürültü"], (D("8.33"), D("0.83"), D("-7.50"), "Kötüleşiyor", "marj daralıyor (son marj %0.83)"))
        self.assertEqual(e["Debi"][2:4], (D("0.50"), "Değişmedi"))
        self.assertEqual(main.karsilastir(l, SPEK, egilim_esik=D("0.4"))["egilim"][2]["yon"], "İyileşiyor")
        self.assertEqual(main.karsilastir(l, SPEK, marj_esik=D("0.5"))["egilim"][1]["dikkat"], "")
        o = next(x for x in r["ozet"] if x["par"] == "Akım" and x["proto"] == "P1")
        self.assertEqual((o["n"], o["ort"], o["min"], o["max"], o["uygunsuz"], o["hedef_fark"]), (2, D("3.0000"), D("2.9"), D("3.1"), 1, D("0.5")))

    def test_prototip_ozeti_ve_uyarilar(self):
        r = main.karsilastir([so("P1", 3.5), so("P1", 61, "Gürültü"), so("P2", 2.5), so("P2", 7, "Bilinmeyen")], SPEK)
        self.assertEqual([(p["proto"], p["uygunsuz"], p["uygunsuz_par"], p["durum"]) for p in r["proto"]], [("P1", 2, ["Akım", "Gürültü"], "Uygunsuzluk var"), ("P2", 0, [], "Tüm sonuçlar uygun")])
        self.assertEqual(sorted(u["tur"] for u in r["uyarilar"]), ["Parametre test edilmemiş", "Parametrenin spesifikasyonu yok"])
        self.assertEqual([(e["par"], e["yon"]) for e in r["egilim"]], [("Akım", "İyileşiyor"), ("Gürültü", "Tek prototip")])
        self.assertIn("son prototipte 1 uygunsuz sonuç", main.karsilastir([so("P1", 2.5), so("P2", 3.5)], SPEK)["egilim"][0]["dikkat"])
        with self.assertRaises(ValueError):
            main.karsilastir([so("P1", 1, "Yok")], SPEK)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "test_sonuclari.csv", ORNEK / "spesifikasyon.csv", cikti)
            self.assertEqual([(p["proto"], p["uygunsuz"], len(p["uygunsuz_par"])) for p in r["proto"]], [("P1", 5, 3), ("P2", 0, 0), ("P3", 1, 1)])
            e = {x["par"]: (x["ilk"], x["son"], x["yon"], bool(x["dikkat"])) for x in r["egilim"]}
            self.assertEqual(e, {"Debi": (D("-3.75"), D("11.25"), "İyileşiyor", False), "Gürültü": (D("-2.50"), D("4.17"), "İyileşiyor", False),
                                 "Akım": (D("0.00"), D("45.00"), "İyileşiyor", False), "Sıcaklık artışı": (D("12.86"), D("0.00"), "Kötüleşiyor", True)})
            self.assertEqual((len(r["sonuclar"]), r["uyarilar"]), (24, []))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Prototip Özeti", "Eğilim", "Parametre Özeti", "Sonuçlar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            s, k = Path(tmp) / "s.csv", Path(tmp) / "k.csv"
            k.write_text("Parametre;Alt Limit;Üst Limit\nA;1,5;2.5\nA;1;2\nB;;\nC;5;5\nD;;10\n", encoding="utf-8")
            s.write_text("Prototip;Parametre;Sonuç\nP1;A;2,0\nP1;D;yok\n;a;1.250\n", encoding="utf-8")
            l, spek, uy = main.oku(s, k)
            self.assertEqual((sorted(spek), spek["a"]["alt"], [(x["proto"], x["deger"]) for x in l]), (["a", "d"], D("1.5"), [("P1", D("2.0")), ("—", D("1.250"))]))
            self.assertEqual(sorted(u["tur"] for u in uy), ["Alt limit üst limitten küçük değil", "Limit tanımı eksik", "Parametre tekrar ediyor", "Sonuç okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--sonuclar", str(s), "--spesifikasyon", str(k), "--cikti", c]
            self.assertEqual(main.main(arg + ["--egilim-esik", "2", "--marj-esik", "5"]), 0)
            self.assertEqual(main.main(arg + ["--marj-esik", "-1"]), 2)
            self.assertEqual(main.main(["--sonuclar", str(Path(tmp) / "yok.csv")]), 1)
            k.write_text("Parametre;Alt Limit;Üst Limit\nB;;\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
