"""
Kestirimci Bakım Uyarıları — Workers / Workless kod bloğu
Üretim › Bakım Onarım › Bakım Mühendisi

Titreşim, sıcaklık ve akım ölçümlerini (portatif cihaz rotası veya sensör dışa aktarımı) ekipman limitleriyle
karşılaştırır ve arıza riski yüksek ekipmanları öne çıkarır:
  - Eşik: son ölçüm alarm (Kritik) veya uyarı (Yüksek) seviyesinde mi?
  - Eğilim: son N gündeki (varsayılan 30) doğrusal eğilimle alarm / uyarı seviyesine kaç gün kaldığı.
  - Ani değişim (titreşim): baz çizgisinden sapma, B/C sınırının %25'inden büyük (ISO 10816 / 20816 "Kriter II").
  - Kayıt / sensör: art arda aynı değer, sıfır / negatif okuma, ölçüm periyodu aşılmış veya hiç ölçüm yok.
  - Aynı ekipmanda birden fazla parametrede belirti.
Titreşim limiti girilmemişse ISO 10816-3 makine grubu / temel tipine göre B/C (uyarı) ve C/D (alarm) sınırları,
akım limiti girilmemişse plaka (nominal) akımı alarm, nominal × 0,90 uyarı olarak kullanılır.
Rapor: ekipman risk sıralaması, bulgular, seri istatistikleri, grafikler. İnternete bağlanmaz.

Kullanım:
    python main.py                                                     # örnek tesis, 90 günlük ölçüm
    python main.py --ekipmanlar ekipmanlar.xlsx --olcumler olcumler.xlsx --pencere 30 --ufuk 30
"""
from __future__ import annotations

import argparse
import csv
import re
import statistics
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")

# ISO 10816-3 (ISO 20816-3) titreşim hızı (mm/s RMS) bölge sınırları: (A/B, B/C, C/D)
# Grup 1: 300 kW – 50 MW büyük makineler; Grup 2: 15 – 300 kW orta büyüklükte makineler.
ISO_10816_3 = {"grup 1 rijit": (2.3, 4.5, 7.1), "grup 1 esnek": (3.5, 7.1, 11.0),
               "grup 2 rijit": (1.4, 2.8, 4.5), "grup 2 esnek": (2.3, 4.5, 7.1)}
PARAMETRELER = {"Titreşim": ("titres", "vib", "hiz rms"), "Sıcaklık": ("sicak", "temp"), "Akım": ("akim", "current", "amper")}
BIRIM = {"Titreşim": "mm/s", "Sıcaklık": "°C", "Akım": "A"}
INCELEME = {"Titreşim": "Spektrum / zarf analiziyle kaynağı belirleyin (dengesizlik, eksen kaçıklığı, rulman, gevşeklik); yatak ve bağlantıları kontrol edin.",
            "Sıcaklık": "Yağlama, yük, soğutma / havalandırma ve yatak durumunu kontrol edin; ortam sıcaklığını not edin.",
            "Akım": "Mekanik yük ve zorlanmayı, faz akımları ile gerilim dengesizliğini ve motor sargılarını kontrol edin."}

EKIPMAN_SUTUNLARI = {"kod": ("ekipman kodu", "ekipman", "makine kodu", "kod"), "ad": ("ekipman adi", "makine adi", "ad", "tanim"),
                     "kritiklik": ("kritiklik", "kritiklik sinifi", "onem"), "iso": ("iso 10816 3 sinifi", "iso sinifi", "iso 10816 sinifi", "makine grubu"),
                     "t_uyari": ("titresim uyari mm s", "titresim uyari"), "t_alarm": ("titresim alarm mm s", "titresim alarm"),
                     "s_uyari": ("sicaklik uyari c", "sicaklik uyari"), "s_alarm": ("sicaklik alarm c", "sicaklik alarm"),
                     "nominal": ("nominal akim a", "nominal akim", "plaka akimi"), "a_uyari": ("akim uyari a", "akim uyari"),
                     "a_alarm": ("akim alarm a", "akim alarm"), "periyot": ("olcum periyodu gun", "olcum periyodu", "periyot gun", "periyot")}
OLCUM_SUTUNLARI = {"tarih": ("tarih", "olcum tarihi", "zaman"), "saat": ("saat",), "kod": ("ekipman kodu", "ekipman", "makine kodu", "kod"),
                   "nokta": ("olcum noktasi", "nokta", "konum"), "parametre": ("parametre", "olcum turu", "tur"), "deger": ("deger", "olcum", "sonuc")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def sayi(x) -> float | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip().replace(",", ".")
    return float(s) if re.fullmatch(r"-?\d+(\.\d+)?", s) else None


def goster(x, ondalik: int = 1) -> str:
    return "—" if x is None else f"{x:.{ondalik}f}".replace(".", ",")


def zaman(t, s) -> datetime | None:
    if isinstance(t, datetime) and not s:
        return t
    d = t.date() if isinstance(t, datetime) else t if isinstance(t, date) else None
    if d is None:
        for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
            try:
                d = datetime.strptime(str(t or "").strip()[:10], f).date()
                break
            except ValueError:
                pass
    if d is None:
        return None
    if isinstance(s, time):
        return datetime.combine(d, s)
    m = re.search(r"(\d{1,2})[:.](\d{2})", str(s or "") or str(t or "")[10:])
    return datetime.combine(d, time(int(m.group(1)), int(m.group(2))) if m else time(0, 0))


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Ekipman:
    kod: str
    ad: str
    kritiklik: str
    iso: str
    limitler: dict          # parametre -> (uyarı, alarm, kaynak)
    periyot: int | None


@dataclass
class Olcum:
    satir: int
    an: datetime
    kod: str
    nokta: str
    parametre: str
    ham: str
    deger: float | None


@dataclass
class Seri:
    kod: str
    nokta: str
    parametre: str
    olcumler: list
    uyari: float | None = None
    alarm: float | None = None
    son: float | None = None
    baz: float | None = None
    egim: float | None = None       # birim / gün
    r2: float | None = None
    gun_uyari: float | None = None
    gun_alarm: float | None = None
    bolge: str = "Normal"
    bulgular: list = field(default_factory=list)

    @property
    def ad(self) -> str:
        return f"{self.kod} · {self.nokta} · {self.parametre}"


def parametre_bul(ham) -> str | None:
    k = katla(ham)
    return next((p for p, on in PARAMETRELER.items() if any(k.startswith(x) or f" {x}" in f" {k}" for x in on)), None)


def ekipman_oku(yol: Path, akim_orani: float = 0.9) -> dict[str, Ekipman]:
    sonuc = {}
    for r in kayitlar(yol, EKIPMAN_SUTUNLARI, ("kod",)):
        kod = str(r.get("kod") or "").strip()
        if not kod:
            continue
        iso = str(r.get("iso") or "").strip()
        lim = {}
        tu, ta = sayi(r.get("t_uyari")), sayi(r.get("t_alarm"))
        if tu is not None or ta is not None:
            lim["Titreşim"] = (tu, ta, "Ekipman tablosu")
        elif katla(iso) in ISO_10816_3:
            _, bc, cd = ISO_10816_3[katla(iso)]
            lim["Titreşim"] = (bc, cd, f"ISO 10816-3 {iso} (B/C, C/D)")
        su, sa = sayi(r.get("s_uyari")), sayi(r.get("s_alarm"))
        if su is not None or sa is not None:
            lim["Sıcaklık"] = (su, sa, "Ekipman tablosu")
        au, aa, nom = sayi(r.get("a_uyari")), sayi(r.get("a_alarm")), sayi(r.get("nominal"))
        if au is not None or aa is not None:
            lim["Akım"] = (au, aa, "Ekipman tablosu")
        elif nom:
            lim["Akım"] = (round(nom * akim_orani, 2), nom, f"Nominal akım {goster(nom)} A (uyarı ×{goster(akim_orani, 2)})")
        p = sayi(r.get("periyot"))
        sonuc[kod] = Ekipman(kod, str(r.get("ad") or "").strip(), str(r.get("kritiklik") or "").strip().upper(), iso, lim, int(p) if p else None)
    return sonuc


def olcum_oku(yol: Path) -> tuple[list[Olcum], list[str]]:
    sonuc, hatalar = [], []
    for r in kayitlar(yol, OLCUM_SUTUNLARI, ("tarih", "kod", "parametre", "deger")):
        if not r.get("kod"):
            continue
        an = zaman(r.get("tarih"), r.get("saat"))
        par = parametre_bul(r.get("parametre"))
        if an is None or par is None:
            hatalar.append(f"Satır {r['_satir']}: " + ("tarih okunamadı" if an is None else f"parametre tanınmadı ('{r.get('parametre')}')"))
            continue
        ham = str(r.get("deger") if r.get("deger") is not None else "").strip()
        sonuc.append(Olcum(r["_satir"], an, str(r["kod"]).strip(), str(r.get("nokta") or "").strip() or "—", par, ham, sayi(r.get("deger"))))
    return sorted(sonuc, key=lambda o: (o.kod, o.nokta, o.parametre, o.an)), hatalar


# ----------------------------------------------------------------------------
# Analiz
# ----------------------------------------------------------------------------

def dogrusal(xs: list[float], ys: list[float]) -> tuple[float, float, float]:
    """En küçük kareler: (eğim, kesişim, R²)."""
    mx, my = statistics.fmean(xs), statistics.fmean(ys)
    sxx = sum((x - mx) ** 2 for x in xs)
    sxy = sum((x - mx) * (y - my) for x, y in zip(xs, ys))
    syy = sum((y - my) ** 2 for y in ys)
    if sxx == 0:
        return 0.0, my, 0.0
    b = sxy / sxx
    return b, my - b * mx, (sxy * sxy / (sxx * syy)) if syy else 0.0


def sabit_beklenmedik(lst: list, ayni: list) -> bool:
    """Seri geri kalanında ölçümden ölçüme belirgin değişiyorsa (medyan |Δ| ≥ 2 × çözünürlük) uzun tekrar beklenmedik sayılır.
    Değişkenliği çözünürlüğün altında kalan serilerde (ör. 1,2 – 1,3 mm/s) tekrar doğaldır."""
    disari = [o.deger for o in lst if o not in ayni and o.deger is not None]
    farklar = [abs(b - a) for a, b in zip(disari, disari[1:])]
    degerler = sorted({o.deger for o in lst if o.deger is not None})
    adim = min((b - a for a, b in zip(degerler, degerler[1:])), default=0)
    if len(farklar) < 3 or adim <= 0:
        return len(ayni) == len(lst)
    return statistics.median(farklar) >= 2 * adim - 1e-9


SIRA = {"Kritik": 0, "Yüksek": 1, "Orta": 2, "Bilgi": 3, "Normal": 4}


def analiz_et(ekipmanlar: dict[str, Ekipman], olcumler: list[Olcum], rapor_tarihi: date | None = None, pencere: int = 30, ufuk: int = 30,
              min_nokta: int = 5, min_r2: float = 0.6, baz_n: int = 5, tekrar_n: int = 6) -> dict:
    bulgular = []

    def bulgu(onem, tur, kod, aciklama, seri=None, parametre=None):
        b = {"onem": onem, "tur": tur, "kod": kod, "seri": seri.ad if seri else "", "parametre": parametre or (seri.parametre if seri else ""),
             "aciklama": aciklama}
        bulgular.append(b)
        if seri:
            seri.bulgular.append(b)

    if rapor_tarihi is None:
        rapor_tarihi = max((o.an.date() for o in olcumler), default=date.today())
    gruplar = defaultdict(list)
    for o in olcumler:
        gruplar[(o.kod, o.nokta, o.parametre)].append(o)
    seriler = []
    for (kod, nokta, par), lst in gruplar.items():
        s = Seri(kod, nokta, par, lst)
        seriler.append(s)
        e = ekipmanlar.get(kod)
        if e and par in e.limitler:
            s.uyari, s.alarm, _ = e.limitler[par]
        birim = BIRIM[par]
        gecerli = [o for o in lst if o.deger is not None and o.deger > 0]
        sifir = [o for o in lst if o.deger is not None and o.deger <= 0]
        okunamayan = [o for o in lst if o.deger is None]
        if sifir:
            bulgu("Bilgi", "Sıfır / negatif değer", kod, f"{s.ad}: {', '.join(f'{o.an:%d.%m.%Y} ({o.ham})' for o in sifir[:5])}. Ekipman duruşta "
                  "mıydı, sensör mü arızalı? Bu değerler eğilim hesabına alınmadı.", s)
        if okunamayan:
            bulgu("Bilgi", "Okunamayan değer", kod, f"{s.ad}: {', '.join(f'{o.an:%d.%m.%Y} (' + repr(o.ham) + ')' for o in okunamayan[:5])}", s)
        if not gecerli:
            continue
        son = gecerli[-1]
        s.son = son.deger
        if s.uyari is None and s.alarm is None:
            bulgu("Bilgi", "Limit tanımsız", kod, f"{s.ad}: ekipman tablosunda {par.lower()} limiti yok; yalnız eğilim ve kayıt kontrolleri yapıldı.", s)
        # Eşik
        if s.alarm is not None and son.deger >= s.alarm:
            s.bolge = "Alarm"
            bulgu("Kritik", "Alarm seviyesi", kod, f"{s.ad}: {son.an:%d.%m.%Y} ölçümü {goster(son.deger)} {birim} ≥ alarm {goster(s.alarm)} {birim}", s)
        elif s.uyari is not None and son.deger >= s.uyari:
            s.bolge = "Uyarı"
            bulgu("Yüksek", "Uyarı seviyesi", kod, f"{s.ad}: {son.an:%d.%m.%Y} ölçümü {goster(son.deger)} {birim} ≥ uyarı {goster(s.uyari)} {birim}"
                  + (f" (alarm {goster(s.alarm)})" if s.alarm is not None else ""), s)
        else:
            gecmis = [o for o in gecerli[:-1] if (s.alarm is not None and o.deger >= s.alarm) or (s.uyari is not None and o.deger >= s.uyari)]
            if gecmis:
                bulgu("Bilgi", "Geçmiş aşım", kod, f"{s.ad}: son ölçüm normal ({goster(son.deger)} {birim}) ama dönemde {len(gecmis)} aşım var: "
                      + ", ".join(f"{o.an:%d.%m.%Y} {goster(o.deger)}" for o in gecmis[:5]) + ". Tekrarlıyorsa nedenini araştırın.", s)
        # Eğilim
        penc = [o for o in gecerli if (son.an - o.an).days <= pencere]
        if len(penc) >= min_nokta:
            xs = [(o.an - penc[0].an).total_seconds() / 86400 for o in penc]
            s.egim, kes, s.r2 = dogrusal(xs, [o.deger for o in penc])
            tahmin = kes + s.egim * xs[-1]
            if s.egim > 0 and s.r2 >= min_r2:
                if s.alarm is not None and son.deger < s.alarm:
                    s.gun_alarm = max((s.alarm - tahmin) / s.egim, 0)
                if s.uyari is not None and son.deger < s.uyari:
                    s.gun_uyari = max((s.uyari - tahmin) / s.egim, 0)
                artis = f"son {pencere} günde {len(penc)} ölçümün eğilimi +{goster(s.egim * 7, 2)} {birim}/hafta (R² {goster(s.r2, 2)})"
                if s.gun_alarm is not None and s.gun_alarm <= ufuk:
                    bulgu("Yüksek", "Eğilim: alarma yaklaşıyor", kod, f"{s.ad}: {artis}; bu hızla alarm seviyesine ({goster(s.alarm)} {birim}) "
                          f"yaklaşık {s.gun_alarm:.0f} gün", s)
                elif s.gun_uyari is not None and s.gun_uyari <= ufuk:
                    bulgu("Orta", "Eğilim: uyarıya yaklaşıyor", kod, f"{s.ad}: {artis}; bu hızla uyarı seviyesine ({goster(s.uyari)} {birim}) "
                          f"yaklaşık {s.gun_uyari:.0f} gün", s)
        # Ani değişim (ISO 10816 / 20816 Kriter II)
        if len(gecerli) > baz_n:
            s.baz = statistics.median(o.deger for o in gecerli[:baz_n])
            if par == "Titreşim" and s.uyari and s.bolge == "Normal" and abs(son.deger - s.baz) > 0.25 * s.uyari:
                bulgu("Orta", "Ani değişim", kod, f"{s.ad}: baz çizgisi {goster(s.baz)} → son {goster(son.deger)} {birim}; değişim "
                      f"{goster(abs(son.deger - s.baz))} > B/C sınırının %25'i ({goster(0.25 * s.uyari, 2)}). Seviye kabul bölgesinde olsa da "
                      "değişimin nedenini araştırın.", s)
        # Sabit değer
        ayni = []
        for o in lst + [None]:
            if o is not None and ayni and o.ham == ayni[-1].ham:
                ayni.append(o)
                continue
            if len(ayni) >= tekrar_n and sabit_beklenmedik(lst, ayni):
                bulgu("Orta", "Sabit değer", kod, f"{s.ad}: {ayni[0].an:%d.%m.%Y} – {ayni[-1].an:%d.%m.%Y} arası art arda {len(ayni)} kez {ayni[0].ham}; "
                      "serinin geri kalanındaki değişkenliğe göre beklenmedik. Sensör donmuş veya değer elle kopyalanmış olabilir; ölçümü doğrulayın.", s)
            ayni = [o] if o is not None else []

    for kod in sorted(set(o.kod for o in olcumler) - set(ekipmanlar)):
        bulgu("Orta", "Tanımsız ekipman", kod, f"{kod}: ölçüm var ama ekipman tablosunda yok; limitler uygulanamadı.")
    son_tarih = {}
    for o in olcumler:
        son_tarih[o.kod] = max(son_tarih.get(o.kod, o.an.date()), o.an.date())
    for kod, e in ekipmanlar.items():
        if kod not in son_tarih:
            bulgu("Orta", "Ölçüm yok", kod, f"{kod} {e.ad}: dönemde hiç ölçüm yok. Kestirimci bakım rotasında mı?")
        elif e.periyot and (rapor_tarihi - son_tarih[kod]).days > e.periyot * 1.5:
            bulgu("Orta", "Ölçüm gecikmiş", kod, f"{kod} {e.ad}: son ölçüm {son_tarih[kod]:%d.%m.%Y}, {(rapor_tarihi - son_tarih[kod]).days} gün önce "
                  f"(periyot {e.periyot} gün)")
    for kod in sorted(set(son_tarih) | set(ekipmanlar)):
        ciddi = {b["parametre"] for b in bulgular if b["kod"] == kod and b["onem"] in ("Kritik", "Yüksek") and b["parametre"]}
        if len(ciddi) >= 2:
            bulgu("Yüksek", "Çoklu belirti", kod, f"{kod}: {' ve '.join(sorted(ciddi))} aynı anda uyarı / alarm veya hızlı eğilim gösteriyor; "
                  "birlikte değerlendirin.")

    bulgular.sort(key=lambda b: (SIRA[b["onem"]], b["kod"]))
    ozet = []
    for kod in sorted(set(son_tarih) | set(ekipmanlar)):
        e = ekipmanlar.get(kod)
        bl = [b for b in bulgular if b["kod"] == kod]
        risk = min((b["onem"] for b in bl), key=SIRA.get, default="Normal")
        pars = sorted({b["parametre"] for b in bl if b["parametre"] and SIRA[b["onem"]] <= SIRA["Orta"]})
        ozet.append({"kod": kod, "ad": e.ad if e else "", "kritiklik": e.kritiklik if e else "", "risk": risk, "bulgu": len(bl),
                     "son_olcum": son_tarih.get(kod), "ozet": "; ".join(dict.fromkeys(b["tur"] for b in bl if b["onem"] != "Bilgi")) or "—",
                     "inceleme": " ".join(INCELEME[p] for p in pars if p in INCELEME)})
    ozet.sort(key=lambda x: (SIRA[x["risk"]], x["kritiklik"] or "Z", x["kod"]))
    return {"bulgular": bulgular, "seriler": sorted(seriler, key=lambda s: (s.kod, s.nokta, s.parametre)), "ozet": ozet, "rapor_tarihi": rapor_tarihi}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Kritik": "FDE2E1", "Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE", "Alarm": "FDE2E1", "Uyarı": "FFF4CE"}
KONTROL = PatternFill("solid", fgColor="FFF4CE")
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def _f(x, n=2):
    return None if x is None else round(x, n)


def rapor_yaz(cikti: Path, ekipmanlar: dict[str, Ekipman], olcumler: list[Olcum], s: dict, hatalar: list[str]) -> None:
    wb = Workbook()
    o = wb.active
    o.title = "Ekipman Riski"
    _baslik(o, ["Sıra", "Ekipman", "Ekipman Adı", "Kritiklik", "Risk", "Bulgu", "Son Ölçüm", "Özet", "Önerilen İnceleme", "Planlanan Müdahale / Karar"],
            (6, 10, 30, 9, 9, 7, 11, 46, 60, 30))
    for i, x in enumerate(s["ozet"], 1):
        o.append([i, x["kod"], x["ad"], x["kritiklik"], x["risk"], x["bulgu"], x["son_olcum"], x["ozet"], x["inceleme"], ""])
        o.cell(o.max_row, 7).number_format = "DD.MM.YYYY"
        if x["risk"] in RENK:
            o.cell(o.max_row, 5).fill = PatternFill("solid", fgColor=RENK[x["risk"]])
        o.cell(o.max_row, 10).fill = KONTROL
        for c in o[o.max_row]:
            c.alignment = UST
    o.append([])
    o.append(["", "Rapor tarihi", s["rapor_tarihi"]])
    o.cell(o.max_row, 3).number_format = "DD.MM.YYYY"
    for h in hatalar:
        o.append(["", "Okunamayan satır", h])
    o.append(["", "Not", "Eğilim tahmini doğrusal varsayıma dayanır; rulman hasarı gibi arızalar son aşamada hızlanabilir. Durdurma / müdahale "
              "kararı bakım mühendisinindir."])

    b = wb.create_sheet("Bulgular")
    _baslik(b, ["Önem", "Tür", "Ekipman", "Seri", "Açıklama", "İnceleme Sonucu"], (9, 24, 9, 40, 90, 28))
    for x in s["bulgular"]:
        b.append([x["onem"], x["tur"], x["kod"], x["seri"], x["aciklama"], ""])
        b.cell(b.max_row, 1).fill = PatternFill("solid", fgColor=RENK[x["onem"]])
        b.cell(b.max_row, 5).alignment = UST
        b.cell(b.max_row, 6).fill = KONTROL

    sr = wb.create_sheet("Seriler")
    _baslik(sr, ["Ekipman", "Ölçüm Noktası", "Parametre", "Birim", "Ölçüm", "Uyarı", "Alarm", "Baz Çizgisi", "Son Değer", "Bölge",
                 "Eğim (birim/hafta)", "R²", "Uyarıya ~Gün", "Alarma ~Gün"], (9, 20, 10, 7, 8, 8, 8, 10, 10, 9, 12, 7, 11, 11))
    for x in s["seriler"]:
        sr.append([x.kod, x.nokta, x.parametre, BIRIM[x.parametre], len(x.olcumler), x.uyari, x.alarm, _f(x.baz), x.son, x.bolge,
                   _f(x.egim * 7 if x.egim is not None else None, 3), _f(x.r2), _f(x.gun_uyari, 0), _f(x.gun_alarm, 0)])
        if x.bolge in RENK:
            sr.cell(sr.max_row, 10).fill = PatternFill("solid", fgColor=RENK[x.bolge])
    sr.auto_filter.ref = f"A1:N{sr.max_row}"

    gr = wb.create_sheet("Grafikler")
    satir = 1
    for x in s["seriler"]:
        if not any(SIRA[bb["onem"]] <= SIRA["Orta"] for bb in x.bulgular):
            continue
        lst = [m for m in x.olcumler if m.deger is not None]
        gr.cell(satir, 1, x.ad).font = Font(bold=True)
        limitler = [(a, v) for a, v in (("Uyarı", x.uyari), ("Alarm", x.alarm)) if v is not None]
        for j, h in enumerate(["Tarih", "Değer"] + [a for a, _ in limitler], 1):
            gr.cell(satir + 1, j, h)
        for i, m in enumerate(lst, satir + 2):
            gr.cell(i, 1, m.an.strftime("%d.%m"))
            gr.cell(i, 2, m.deger)
            for j, (_, v) in enumerate(limitler, 3):
                gr.cell(i, j, v)
        ch = LineChart()
        ch.title, ch.height, ch.width = x.ad, 7, 22
        ch.y_axis.title = BIRIM[x.parametre]
        ch.add_data(Reference(gr, min_col=2, max_col=2 + len(limitler), min_row=satir + 1, max_row=satir + 1 + len(lst)), titles_from_data=True)
        ch.set_categories(Reference(gr, min_col=1, min_row=satir + 2, max_row=satir + 1 + len(lst)))
        gr.add_chart(ch, f"F{satir}")
        satir += max(len(lst) + 4, 17)

    ol = wb.create_sheet("Ölçümler")
    _baslik(ol, ["Satır", "Tarih", "Ekipman", "Ölçüm Noktası", "Parametre", "Değer"], (6, 16, 9, 20, 10, 9))
    for m in sorted(olcumler, key=lambda m: (m.an, m.kod, m.parametre)):
        ol.append([m.satir, m.an, m.kod, m.nokta, m.parametre, m.deger if m.deger is not None else m.ham])
        ol.cell(ol.max_row, 2).number_format = "DD.MM.YYYY HH:MM" if m.an.time() != time(0, 0) else "DD.MM.YYYY"
    ol.auto_filter.ref = f"A1:F{ol.max_row}"

    ek = wb.create_sheet("Ekipmanlar")
    _baslik(ek, ["Ekipman", "Ad", "Kritiklik", "ISO 10816-3 Sınıfı", "Parametre", "Uyarı", "Alarm", "Limit Kaynağı", "Ölçüm Periyodu (gün)"],
            (9, 30, 9, 16, 10, 8, 8, 40, 12))
    for e in ekipmanlar.values():
        for par in ("Titreşim", "Sıcaklık", "Akım"):
            u, a, kaynak = e.limitler.get(par, (None, None, "Tanımsız"))
            ek.append([e.kod, e.ad, e.kritiklik, e.iso, par, u, a, kaynak, e.periyot])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(ekipman_yolu: Path, olcum_yolu: Path, cikti: Path, rapor_tarihi: date | None = None, pencere: int = 30, ufuk: int = 30,
             akim_orani: float = 0.9, tekrar_n: int = 6) -> dict:
    ekipmanlar = ekipman_oku(ekipman_yolu, akim_orani)
    olcumler, hatalar = olcum_oku(olcum_yolu)
    s = analiz_et(ekipmanlar, olcumler, rapor_tarihi, pencere, ufuk, tekrar_n=tekrar_n)
    rapor_yaz(cikti, ekipmanlar, olcumler, s, hatalar)
    return {**s, "ekipmanlar": ekipmanlar, "olcumler": olcumler, "hatalar": hatalar}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Titreşim, sıcaklık ve akım ölçümlerinde eşik ve eğilim analiziyle arıza riski yüksek ekipmanları sıralar.")
    p.add_argument("--ekipmanlar", type=Path, default=ORNEK / "ekipmanlar.csv", help="Ekipman listesi: kod, kritiklik, ISO sınıfı, limitler, periyot")
    p.add_argument("--olcumler", type=Path, default=ORNEK / "olcumler.csv", help="Ölçümler: tarih, ekipman, nokta, parametre, değer")
    p.add_argument("--rapor-tarihi", help="GG.AA.YYYY (varsayılan: en son ölçüm tarihi)")
    p.add_argument("--pencere", type=int, default=30, help="Eğilim için geriye bakılan gün (varsayılan 30)")
    p.add_argument("--ufuk", type=int, default=30, help="Bu kadar gün içinde alarm / uyarıya ulaşacaksa bildir (varsayılan 30)")
    p.add_argument("--akim-uyari-orani", type=float, default=0.9, help="Akım limiti yoksa uyarı = nominal × oran (varsayılan 0,90)")
    p.add_argument("--tekrar", type=int, default=6, help="Art arda kaç aynı değer şüpheli sayılır (varsayılan 6)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kestirimci_bakim_uyarilari.xlsx")
    a = p.parse_args(argv)
    for y in (a.ekipmanlar, a.olcumler):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    rt = None
    if a.rapor_tarihi:
        try:
            rt = datetime.strptime(a.rapor_tarihi, "%d.%m.%Y").date()
        except ValueError:
            print("[X] --rapor-tarihi GG.AA.YYYY biçiminde olmalı")
            return 1
    try:
        s = calistir(a.ekipmanlar, a.olcumler, a.cikti, rt, a.pencere, a.ufuk, a.akim_uyari_orani, a.tekrar)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(s['olcumler'])} ölçüm · {len(s['ekipmanlar'])} ekipman · {len(s['bulgular'])} bulgu (rapor tarihi {s['rapor_tarihi']:%d.%m.%Y})")
    for x in s["ozet"]:
        if x["risk"] in ("Kritik", "Yüksek", "Orta"):
            print(f"[{'X' if x['risk'] == 'Kritik' else '!'}] {x['risk']:6} {x['kod']} {x['ad']}: {x['ozet']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
