import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "k.xlsx"
        cls.s = main.calistir(ORNEK / "ekipmanlar.csv", ORNEK / "olcumler.csv", cls.cikti)
        cls.b = {}
        for b in cls.s["bulgular"]:
            cls.b.setdefault(b["tur"], []).append(b)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_limitler(self):
        e = self.s["ekipmanlar"]
        self.assertEqual(e["P-101"].limitler["Titreşim"][:2], (2.8, 4.5))
        self.assertEqual(e["F-201"].limitler["Titreşim"][:2], (4.5, 7.1))
        self.assertEqual(e["G-601"].limitler["Titreşim"][:2], (4.5, 7.1))
        self.assertEqual(e["M-401"].limitler["Titreşim"], (2.8, 4.5, "Ekipman tablosu"))
        self.assertEqual(e["M-401"].limitler["Akım"][:2], (25.2, 28.0))
        self.assertEqual(self.s["rapor_tarihi"], date(2026, 9, 26))

    def test_risk_sirasi(self):
        self.assertEqual([(x["kod"], x["risk"]) for x in self.s["ozet"]],
                         [("F-201", "Kritik"), ("P-101", "Yüksek"), ("M-401", "Yüksek"), ("K-301", "Orta"), ("P-102", "Orta"),
                          ("H-701", "Orta"), ("R-501", "Orta"), ("G-601", "Bilgi")])

    def test_esik_ve_egilim(self):
        self.assertEqual([b["seri"] for b in self.b["Alarm seviyesi"]], ["F-201 · Fan yatağı · Titreşim"])
        self.assertEqual(sorted(b["kod"] for b in self.b["Uyarı seviyesi"]), ["F-201", "M-401", "P-101"])
        egilim = {b["kod"]: b["aciklama"] for b in self.b["Eğilim: alarma yaklaşıyor"]}
        self.assertEqual(sorted(egilim), ["F-201", "M-401", "P-101"])
        self.assertIn("yaklaşık 13 gün", egilim["P-101"])
        self.assertIn("yaklaşık 21 gün", egilim["M-401"])
        self.assertEqual([b["kod"] for b in self.b["Çoklu belirti"]], ["F-201"])

    def test_diger_bulgular(self):
        self.assertIn("baz çizgisi 1,2 → son 2,1", self.b["Ani değişim"][0]["aciklama"])
        self.assertEqual([b["seri"] for b in self.b["Sabit değer"]], ["P-102 · Motor DE yatak · Sıcaklık"])
        self.assertIn("art arda 10 kez 58,0", self.b["Sabit değer"][0]["aciklama"])
        self.assertIn("15 gün önce", self.b["Ölçüm gecikmiş"][0]["aciklama"])
        self.assertEqual(self.b["Ölçüm yok"][0]["kod"], "H-701")
        self.assertIn("09.08.2026 7,4", self.b["Geçmiş aşım"][0]["aciklama"])
        self.assertEqual(self.b["Sıfır / negatif değer"][0]["kod"], "P-101")

    def test_excel(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Ekipman Riski", "Bulgular", "Seriler", "Grafikler", "Ölçümler", "Ekipmanlar"])
        self.assertEqual(wb["Ekipman Riski"]["B2"].value, "F-201")
        self.assertEqual(wb["Ekipman Riski"]["J1"].value, "Planlanan Müdahale / Karar")
        self.assertEqual(len(wb["Grafikler"]._charts), 6)


class KuralTesti(unittest.TestCase):
    EKIPMAN = "Ekipman Kodu;Sıcaklık Uyarı;Sıcaklık Alarm\nX;80;90\n"

    def yaz(self, t, ekipman, olcum):
        (t / "e.csv").write_text(ekipman, encoding="utf-8")
        (t / "o.csv").write_text(olcum, encoding="utf-8")
        return main.calistir(t / "e.csv", t / "o.csv", t / "x.xlsx")

    def seri(self, degerler):
        return "Tarih;Ekipman Kodu;Ölçüm Noktası;Parametre;Değer\n" + "".join(
            f"{g:02d}.09.2026;X;Yatak;Sıcaklık;{v}\n" for g, v in enumerate(degerler, 1))

    def test_sabit_deger_cozunurluk(self):
        with tempfile.TemporaryDirectory() as tmp:
            s = self.yaz(Path(tmp), self.EKIPMAN, self.seri(["50,1", "52,4", "49,8", "51,6", "53,0"] + ["55,0"] * 6 + ["50,9"]))
            self.assertEqual([b["tur"] for b in s["bulgular"]], ["Sabit değer"])
            s = self.yaz(Path(tmp), self.EKIPMAN, self.seri(["55,1", "55,0", "55,1", "55,0", "54,9"] + ["55,0"] * 6 + ["55,1"]))
            self.assertEqual(s["bulgular"], [])

    def test_tanimsiz_ve_limitsiz(self):
        with tempfile.TemporaryDirectory() as tmp:
            s = self.yaz(Path(tmp), "Ekipman Kodu;Ekipman Adı\nA;Pompa\n",
                         "Tarih;Ekipman Kodu;Parametre;Değer\n01.09.2026;A;Titreşim hızı RMS;1,1\n01.09.2026;Z;Akım;10\n01.09.2026;A;Basınç;3\n")
            turler = {(b["tur"], b["kod"]) for b in s["bulgular"]}
            self.assertIn(("Limit tanımsız", "A"), turler)
            self.assertIn(("Tanımsız ekipman", "Z"), turler)
            self.assertIn("parametre tanınmadı", s["hatalar"][0])

    def test_parametreler(self):
        with tempfile.TemporaryDirectory() as tmp:
            s = main.calistir(ORNEK / "ekipmanlar.csv", ORNEK / "olcumler.csv", Path(tmp) / "x.xlsx", ufuk=10)
            self.assertNotIn("Eğilim: alarma yaklaşıyor", {b["tur"] for b in s["bulgular"]})

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--olcumler", str(Path(tmp) / "yok.csv")]), 1)
            self.assertEqual(main.main(["--rapor-tarihi", "2026-09-30", "--cikti", str(Path(tmp) / "x.xlsx")]), 1)


if __name__ == "__main__":
    unittest.main()
