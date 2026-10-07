"""
Kumaş Kontrol (4 Puan Sistemi) Raporu — Workers / Workless kod bloğu
Tekstil › Kalite Kontrol

Top bazında kumaş hata kayıtlarından 4 puan sistemine (ASTM D5430'da tanımlanan yaygın uygulama) göre puan
hesaplar, 100 yard² (ve 100 m²) başına puanı bulur, topu ve partiyi kabul/ret olarak raporlar.

Puanlama (hata uzunluğu, kumaş boyu veya eni yönünde):
    ≤ 3 inç (7,62 cm) → 1 · 3–6 inç (15,24 cm) → 2 · 6–9 inç (22,86 cm) → 3 · > 9 inç → 4 puan
    Delik / açıklık: ≤ 1 inç (2,54 cm) → 2 · > 1 inç → 4 puan
    Bir yardda en fazla 4 puan verilir. 1 yarddan uzun (sürekli) hata, uzandığı her yard için 4 puan alır.
Formül:
    Puan / 100 yd² = toplam puan × 36 × 100 / (kesilebilir en [inç] × top uzunluğu [yd])
Kabul: varsayılan top ≤ 40 puan / 100 yd² (alıcı kriterine göre --top-esik ile değiştirin). İnternete bağlanmaz.

Kullanım:
    python main.py                                          # örnek toplar ve hata kayıtlarıyla dener
    python main.py --toplar toplar.xlsx --hatalar hatalar.xlsx
    python main.py --toplar toplar.xlsx --hatalar hatalar.xlsx --birim inc --top-esik 28 --parti-esik 20
"""
from __future__ import annotations

import argparse
import csv
import math
import sys
from collections import Counter, OrderedDict, defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
INC_CM = 2.54
YARD_M = 0.9144


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def sayi(x, varsayilan: float | None = None) -> float | None:
    if x in (None, ""):
        return varsayilan
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip().replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(Decimal(s))
    except InvalidOperation:
        return varsayilan


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def bul(baslik, *adlar):
    b = [kucuk(x) for x in baslik]
    return next((i for i, x in enumerate(b) if x in adlar), None)


# ----------------------------------------------------------------------------
# Puanlama
# ----------------------------------------------------------------------------

def hata_puani(uzunluk_inc: float, delik: bool) -> int:
    if delik:
        return 2 if uzunluk_inc <= 1 else 4
    if uzunluk_inc <= 3:
        return 1
    if uzunluk_inc <= 6:
        return 2
    if uzunluk_inc <= 9:
        return 3
    return 4


def top_puani(hatalar: list[dict]) -> tuple[int, dict[int, int]]:
    """Yard bazında puanlar (her yard en fazla 4). Sürekli hata, uzandığı her yarda 4 puan yazar.
    hatalar: {konum_yd, uzunluk_inc, delik, yon} — yon 'boy' ise uzunluk yard boyunca ilerler."""
    yardlar: dict[int, int] = defaultdict(int)
    for h in hatalar:
        ilk = int(math.floor(h["konum_yd"]))
        if h.get("puan") is not None:
            yardlar[ilk] += h["puan"]
            continue
        if not h["delik"] and h.get("yon") == "boy" and h["uzunluk_inc"] > 36:
            kapsanan = math.ceil(h["uzunluk_inc"] / 36 - 1e-9)
            for i in range(kapsanan):
                yardlar[ilk + i] += 4
        else:
            yardlar[ilk] += hata_puani(h["uzunluk_inc"], h["delik"])
    sinirli = {y: min(4, p) for y, p in yardlar.items()}
    return sum(sinirli.values()), sinirli


def puan_100yd2(puan: float, en_inc: float, boy_yd: float) -> float | None:
    return None if not en_inc or not boy_yd else puan * 36 * 100 / (en_inc * boy_yd)


def puan_100m2(puan: float, en_cm: float, boy_m: float) -> float | None:
    return None if not en_cm or not boy_m else puan * 100 / (en_cm / 100 * boy_m)


# ----------------------------------------------------------------------------
# Okuma ve hesap
# ----------------------------------------------------------------------------

def calistir(toplar_yolu: Path, hatalar_yolu: Path, cikti: Path, birim: str = "metrik", top_esik: float = 40,
             parti_esik: float | None = None, en_tolerans: float = 2.0) -> dict:
    parti_esik = top_esik if parti_esik is None else parti_esik
    uz_carpan, konum_carpan, en_carpan = (1 / INC_CM, 1 / YARD_M, 1 / INC_CM) if birim == "metrik" else (1.0, 1.0, 1.0)
    # metrik: uzunluk cm → inç, konum m → yd, en cm → inç

    t = tablo_oku(toplar_yolu)
    i_no = bul(t[0], "top no", "top", "rulo no", "top numarası")
    i_parti = bul(t[0], "parti", "parti no", "lot", "sipariş", "sipariş no")
    i_boy = bul(t[0], "kontrol edilen uzunluk", "ölçülen uzunluk", "uzunluk", "boy", "metre", "yard")
    i_etiket_boy = bul(t[0], "etiket uzunluğu", "etiket boyu", "fatura uzunluğu")
    i_en = bul(t[0], "kesilebilir en", "ölçülen en", "en")
    i_etiket_en = bul(t[0], "sipariş eni", "etiket eni", "istenen en")
    i_renk = bul(t[0], "renk", "renk kodu")
    if None in (i_no, i_boy, i_en):
        raise SystemExit(f"Top listesinde Top No, Uzunluk ve Kesilebilir En gerekli. Başlıklar: {t[0]}")
    toplar: "OrderedDict[str, dict]" = OrderedDict()
    for r in t[1:]:
        no = str(r[i_no] or "").strip()
        if no:
            toplar[no] = {"no": no, "parti": str(r[i_parti] or "").strip() if i_parti is not None else "",
                          "renk": str(r[i_renk] or "").strip() if i_renk is not None else "",
                          "boy": sayi(r[i_boy], 0.0), "en": sayi(r[i_en], 0.0),
                          "etiket_boy": sayi(r[i_etiket_boy]) if i_etiket_boy is not None else None,
                          "etiket_en": sayi(r[i_etiket_en]) if i_etiket_en is not None else None, "hatalar": []}

    h = tablo_oku(hatalar_yolu)
    j_no = bul(h[0], "top no", "top", "rulo no", "top numarası")
    j_konum = bul(h[0], "konum", "metre", "yard", "konum (m)", "konum (yd)")
    j_tur = bul(h[0], "hata türü", "hata", "hata kodu", "kusur")
    j_uz = bul(h[0], "uzunluk", "hata uzunluğu", "uzunluk (cm)", "uzunluk (inç)")
    j_delik = bul(h[0], "delik", "delik/açıklık", "açıklık")
    j_yon = bul(h[0], "yön", "doğrultu")
    j_puan = bul(h[0], "puan")
    if None in (j_no, j_konum) or (j_uz is None and j_puan is None):
        raise SystemExit(f"Hata kayıtlarında Top No, Konum ve Uzunluk (veya Puan) gerekli. Başlıklar: {h[0]}")
    uyarilar = []
    for r in h[1:]:
        no = str(r[j_no] or "").strip()
        if not no:
            continue
        if no not in toplar:
            uyarilar.append(f"Hata kaydındaki {no} topu top listesinde yok; atlandı")
            continue
        tur = str(r[j_tur] or "").strip() if j_tur is not None else ""
        delik = (kucuk(r[j_delik]) in ("e", "evet", "x", "1", "var")) if j_delik is not None else ("delik" in kucuk(tur) or "yırtık" in kucuk(tur))
        yon = kucuk(r[j_yon]) if j_yon is not None else ""
        yon = "boy" if yon.startswith(("boy", "çözgü", "uzun")) else "en" if yon else ""
        puan = sayi(r[j_puan]) if j_puan is not None else None
        toplar[no]["hatalar"].append({"konum_yd": sayi(r[j_konum], 0.0) * konum_carpan,
                                      "uzunluk_inc": (sayi(r[j_uz], 0.0) if j_uz is not None else 0.0) * uz_carpan,
                                      "delik": delik, "yon": yon, "tur": tur or "Belirtilmemiş",
                                      "puan": int(puan) if puan is not None else None})

    for top in toplar.values():
        boy_yd = top["boy"] * konum_carpan
        en_inc = top["en"] * en_carpan
        top["puan"], top["yard_puan"] = top_puani(top["hatalar"])
        top["p100yd2"] = puan_100yd2(top["puan"], en_inc, boy_yd)
        en_cm = top["en"] if birim == "metrik" else top["en"] * INC_CM
        boy_m = top["boy"] if birim == "metrik" else top["boy"] * YARD_M
        top["p100m2"] = puan_100m2(top["puan"], en_cm, boy_m)
        top["boy_yd"], top["en_inc"] = boy_yd, en_inc
        notlar = []
        if top["p100yd2"] is None:
            top["sonuc"] = "Ölçü eksik"
        else:
            top["sonuc"] = "KABUL" if top["p100yd2"] <= top_esik else "RET"
        if top["etiket_boy"] and top["boy"] and top["boy"] < top["etiket_boy"]:
            notlar.append(f"ölçülen uzunluk etiketin {top['etiket_boy'] - top['boy']:g} altında")
        if top["etiket_en"] and top["en"] and top["en"] < top["etiket_en"] - en_tolerans:
            notlar.append(f"kesilebilir en istenenin {top['etiket_en'] - top['en']:g} altında")
        dolu = sorted(y for y, p in top["yard_puan"].items() if p >= 4)
        ardisik = max((len(list(g)) for g in _ardisik(dolu)), default=0)
        if ardisik >= 3:
            notlar.append(f"{ardisik} ardışık yardda 4'er puan (sürekli hata / kesim planında dikkat)")
        top["notlar"] = notlar

    partiler: "OrderedDict[str, dict]" = OrderedDict()
    for top in toplar.values():
        p = partiler.setdefault(top["parti"] or "-", {"toplar": [], "puan": 0, "alan_yd2": 0.0})
        p["toplar"].append(top["no"])
        p["puan"] += top["puan"]
        p["alan_yd2"] += top["boy_yd"] * top["en_inc"] / 36
        p["ret"] = p.get("ret", 0) + (top["sonuc"] == "RET")
    for p in partiler.values():
        p["p100yd2"] = p["puan"] * 100 / p["alan_yd2"] if p["alan_yd2"] else None
        p["sonuc"] = "KABUL" if p["p100yd2"] is not None and p["p100yd2"] <= parti_esik else "RET"
    _rapor(toplar, partiler, cikti, birim, top_esik, parti_esik, uyarilar)
    return {"toplar": toplar, "partiler": partiler, "uyarilar": uyarilar}


def _ardisik(sirali: list[int]):
    grup = []
    for y in sirali:
        if grup and y != grup[-1] + 1:
            yield grup
            grup = []
        grup.append(y)
    if grup:
        yield grup


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
YESIL = PatternFill("solid", fgColor="E3F5E1")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def _rapor(toplar, partiler, cikti, birim, top_esik, parti_esik, uyarilar):
    b_boy, b_en = ("m", "cm") if birim == "metrik" else ("yd", "inç")
    wb = Workbook()
    ws = wb.active
    ws.title = "Toplar"
    _baslik(ws, ["Top No", "Parti", "Renk", f"Uzunluk ({b_boy})", f"Kesilebilir En ({b_en})", "Hata Sayısı", "Toplam Puan",
                 "Puan / 100 yd²", "Puan / 100 m²", "Sonuç", "Notlar"], (12, 12, 12, 12, 16, 11, 11, 13, 13, 9, 60))
    for t in toplar.values():
        ws.append([t["no"], t["parti"], t["renk"], t["boy"], t["en"], len(t["hatalar"]), t["puan"],
                   None if t["p100yd2"] is None else round(t["p100yd2"], 2), None if t["p100m2"] is None else round(t["p100m2"], 2),
                   t["sonuc"], "; ".join(t["notlar"])])
        ws.cell(ws.max_row, 10).fill = YESIL if t["sonuc"] == "KABUL" else KIRMIZI
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    n = len(toplar)
    if n:
        g = BarChart()
        g.title, g.height, g.width = f"Top bazında puan / 100 yd² (kabul sınırı {top_esik:g})", 8, 20
        g.add_data(Reference(ws, min_col=8, min_row=1, max_row=1 + n), titles_from_data=True)
        g.set_categories(Reference(ws, min_col=1, min_row=2, max_row=1 + n))
        ws.add_chart(g, "M2")

    p = wb.create_sheet("Partiler", 0)
    p.append([f"Kumaş kontrol raporu — 4 puan sistemi · top kabul ≤ {top_esik:g}, parti kabul ≤ {parti_esik:g} puan / 100 yd²"])
    p["A1"].font = Font(bold=True, size=12)
    p.append([])
    _baslik(p, ["Parti", "Top Sayısı", "Reddedilen Top", "Toplam Puan", "Alan (yd²)", "Parti Puanı / 100 yd²", "Sonuç"],
            (14, 11, 15, 12, 12, 20, 10))
    for ad, x in partiler.items():
        p.append([ad, len(x["toplar"]), x["ret"], x["puan"], round(x["alan_yd2"], 1),
                  None if x["p100yd2"] is None else round(x["p100yd2"], 2), x["sonuc"]])
        p.cell(p.max_row, 7).fill = YESIL if x["sonuc"] == "KABUL" else KIRMIZI
    p.append([])
    p.append(["Parti puanı, toplam puanın toplam alana oranıdır (alan ağırlıklı); kabul edilen bir partide reddedilen tek tek toplar ayrıca değerlendirilir."])
    for u in uyarilar:
        p.append(["Uyarı", u])

    h = wb.create_sheet("Hata Analizi")
    _baslik(h, ["Hata Türü", "Adet", "Pay"], (30, 10, 10))
    sayac = Counter(x["tur"] for t in toplar.values() for x in t["hatalar"])
    toplam = sum(sayac.values())
    for tur, adet in sayac.most_common():
        h.append([tur, adet, adet / toplam if toplam else None])
        h.cell(h.max_row, 3).number_format = "0.0%"
    if sayac:
        g = BarChart()
        g.title, g.height, g.width = "Hata türleri (Pareto)", 8, 18
        g.add_data(Reference(h, min_col=2, min_row=1, max_row=1 + len(sayac)), titles_from_data=True)
        g.set_categories(Reference(h, min_col=1, min_row=2, max_row=1 + len(sayac)))
        h.add_chart(g, "E2")

    d = wb.create_sheet("Hata Kayıtları")
    _baslik(d, ["Top No", "Konum (yd)", "Hata Türü", "Uzunluk (inç)", "Delik", "Yön", "Hata Puanı"], (12, 11, 26, 13, 7, 7, 11))
    for t in toplar.values():
        for x in sorted(t["hatalar"], key=lambda x: x["konum_yd"]):
            pu = x["puan"] if x["puan"] is not None else (
                4 * math.ceil(x["uzunluk_inc"] / 36 - 1e-9) if (x["yon"] == "boy" and not x["delik"] and x["uzunluk_inc"] > 36)
                else hata_puani(x["uzunluk_inc"], x["delik"]))
            d.append([t["no"], round(x["konum_yd"], 2), x["tur"], round(x["uzunluk_inc"], 2), "E" if x["delik"] else "", x["yon"], pu])
    d.freeze_panes = "A2"
    d.auto_filter.ref = d.dimensions

    b = wb.create_sheet("Bilgi")
    for s in [["Hata puanı", "≤ 3 inç (7,62 cm) 1 · 3–6 inç 2 · 6–9 inç 3 · > 9 inç 4; delik/açıklık ≤ 1 inç 2, > 1 inç 4"],
              ["Yard sınırı", "Bir yarda (0,9144 m) en fazla 4 puan; 1 yarddan uzun boyuna (sürekli) hata uzandığı her yarda 4 puan"],
              ["100 yd²", "Toplam puan × 36 × 100 / (kesilebilir en [inç] × uzunluk [yd])"],
              ["100 m²", "Toplam puan × 100 / (kesilebilir en [m] × uzunluk [m]) — 100 yd² = 83,61 m² olduğundan değer farklıdır"],
              ["Kabul", f"Top ≤ {top_esik:g}, parti ≤ {parti_esik:g} puan / 100 yd²; kabul kriterleri alıcıya göre değişir"],
              ["Hata Kayıtları", "Hata puanı sütunu yard sınırı uygulanmadan önceki puandır; top puanında sınır uygulanır"]]:
        b.append(s)
    b.column_dimensions["A"].width = 16
    b.column_dimensions["B"].width = 120
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description="4 puan sistemine göre top ve parti bazında kumaş kontrol raporu.")
    ap.add_argument("--toplar", type=Path, default=BURASI / "ornek_veri" / "toplar.csv",
                    help="Top listesi (.xlsx/.csv): Top No, Uzunluk, Kesilebilir En; Parti, Renk, Etiket Uzunluğu, Sipariş Eni isteğe bağlı")
    ap.add_argument("--hatalar", type=Path, default=BURASI / "ornek_veri" / "hata_kayitlari.csv",
                    help="Hata kayıtları (.xlsx/.csv): Top No, Konum, Uzunluk (veya Puan); Hata Türü, Delik (E/H), Yön (boy/en) isteğe bağlı")
    ap.add_argument("--birim", choices=["metrik", "inc"], default="metrik",
                    help="metrik: uzunluk/en m-cm, hata cm (varsayılan) · inc: yd ve inç")
    ap.add_argument("--top-esik", type=float, default=40, help="Top kabul sınırı, puan / 100 yd² (varsayılan 40)")
    ap.add_argument("--parti-esik", type=float, help="Parti kabul sınırı (varsayılan top sınırıyla aynı)")
    ap.add_argument("--en-tolerans", type=float, default=2.0, help="Kesilebilir en toleransı (cm veya inç, varsayılan 2)")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "kumas_kontrol.xlsx")
    a = ap.parse_args(argv)
    s = calistir(a.toplar, a.hatalar, a.cikti, a.birim, a.top_esik, a.parti_esik, a.en_tolerans)
    ret = [t for t in s["toplar"].values() if t["sonuc"] == "RET"]
    print(f"[OK] {len(s['toplar'])} top · {len(s['partiler'])} parti · reddedilen top: {len(ret)}")
    for ad, p in s["partiler"].items():
        print(f"     Parti {ad}: {p['p100yd2']:.2f} puan / 100 yd² → {p['sonuc']}".replace(".", ","))
    for t in ret:
        print(f"[!] Top {t['no']}: {t['p100yd2']:.2f} puan / 100 yd² → RET".replace(".", ","))
    for u in s["uyarilar"]:
        print(f"[!] {u}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
