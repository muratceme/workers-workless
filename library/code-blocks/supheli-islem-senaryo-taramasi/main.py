"""
Şüpheli İşlem Senaryo Taraması — Workers / Workless kod bloğu
Bankacılık › Kurumsal Uyum › Uyum Uzmanı

İşlem verisini yapılandırılabilir senaryolarla tarar ve uyum biriminin inceleyeceği bir alarm listesi çıkarır.
Bir işlemin şüpheli olup olmadığına KARAR VERMEZ; yalnızca incelenmesi gereken örüntüleri öne çıkarır.

Senaryolar (eşikler ve ağırlıklar senaryolar.json'dan; hiçbir yasal eşik kodun içine gömülmemiştir):
  S1 Parçalı işlem       — eşiğin hemen altında, kısa sürede tekrarlanan nakit/transfer işlemleri
  S2 Hızlı geçiş         — gelen fonun kısa sürede büyük oranda çıkması (geçiş hesabı)
  S3 Nakit yoğunluğu     — dönem girişlerinde nakdin payı ve tutarı yüksek
  S4 Riskli ülke         — kullanıcının verdiği ülke listesine/listesinden transfer
  S5 Çok sayıda gönderen — kısa sürede çok sayıda farklı kişiden para girişi (hesap kiralama örüntüsü)
  S6 Profil dışı tutar   — beyan gelirine veya müşterinin kendi geçmişine göre olağandışı işlem
  S7 Uyuyan hesap        — uzun süre hareketsiz hesapta ani yüksek hareket
  S8 Açıklama anahtar kelimesi — kullanıcının verdiği kelimeler (ör. yasa dışı bahis terimleri)
  S9 Yuvarlak tutar tekrarı
İnternete bağlanmaz.

Kullanım:
    python main.py                                         # örnek verilerle dener
    python main.py --islemler islemler.xlsx --musteriler musteriler.xlsx --senaryolar senaryolar.json
    python main.py --islemler islemler.csv --senaryolar senaryolar.json --riskli-ulkeler ulkeler.csv --kur USD=41,20
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from statistics import median

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
SIFIR = Decimal(0)

TUR_ESLEME = {  # işlem tipi metnindeki ifade → (kanal, yön)
    "nakit yat": ("nakit", 1), "para yat": ("nakit", 1), "nakit çek": ("nakit", -1), "para çek": ("nakit", -1),
    "gelen": ("transfer", 1), "giden": ("transfer", -1), "swift gelen": ("swift", 1), "swift giden": ("swift", -1),
    "kart": ("kart", -1), "pos": ("kart", -1), "atm yat": ("nakit", 1), "atm çek": ("nakit", -1),
}


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def katla(s) -> str:
    s = kucuk(s)
    for a, b in zip("ıöüşğç", "iousgc"):
        s = s.replace(a, b)
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def sayi(x) -> Decimal:
    if x in (None, ""):
        return SIFIR
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("₺", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1:
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return SIFIR


def tl(x) -> str:
    return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def zaman(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    s = str(x or "").strip()
    for b in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%d.%m.%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, b)
        except ValueError:
            continue
    return None


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def _bul(baslik, *adlar):
    b = [katla(x) for x in baslik]
    return next((b.index(katla(a)) for a in adlar if katla(a) in b), None)


def _al(r, i):
    return r[i] if i is not None and i < len(r) else None


@dataclass
class Islem:
    no: str
    zaman: datetime
    musteri: str
    tip: str
    kanal: str
    yon: int              # +1 giriş, −1 çıkış
    tutar: Decimal        # TL, pozitif
    karsi: str
    ulke: str
    aciklama: str


# ----------------------------------------------------------------------------
# Okuma
# ----------------------------------------------------------------------------

def islemleri_oku(yol: Path, kurlar: dict[str, Decimal]) -> tuple[list[Islem], list[str]]:
    s = tablo_oku(yol)
    b = s[0]
    i = {k: _bul(b, *v) for k, v in {
        "no": ("işlem no", "işlem numarası", "referans", "dekont no", "id"),
        "zaman": ("tarih", "işlem tarihi", "tarih saat", "zaman"),
        "saat": ("saat", "işlem saati"),
        "musteri": ("müşteri no", "müşteri numarası", "müşteri", "hesap no", "hesap numarası"),
        "tip": ("işlem tipi", "işlem türü", "tip", "tür", "kanal"),
        "tutar": ("tutar", "işlem tutarı", "miktar"),
        "pb": ("para birimi", "döviz", "pb"),
        "yon": ("yön", "borç alacak", "b a"),
        "karsi": ("karşı taraf", "karşı taraf adı", "gönderen alıcı", "karşı hesap", "karşı taraf iban"),
        "ulke": ("karşı taraf ülke", "ülke", "ülke kodu"),
        "aciklama": ("açıklama", "işlem açıklaması", "not"),
    }.items()}
    for zorunlu, ad in (("zaman", "Tarih"), ("musteri", "Müşteri No"), ("tutar", "Tutar"), ("tip", "İşlem Tipi")):
        if i[zorunlu] is None:
            raise SystemExit(f"İşlem dosyasında '{ad}' sütunu bulunamadı. Başlıklar: {b}")
    islemler, uyarilar, kursuz = [], [], set()
    for n, r in enumerate(s[1:], 2):
        t = zaman(_al(r, i["zaman"]))
        if t and i["saat"] is not None and _al(r, i["saat"]):
            st = zaman(f"{t:%d.%m.%Y} {str(_al(r, i['saat'])).strip()[:8]}")
            t = st or t
        musteri = str(_al(r, i["musteri"]) or "").strip()
        if not t or not musteri:
            continue
        tip = str(_al(r, i["tip"]) or "").strip()
        kt = kucuk(tip)
        kanal, yon = next((v for k, v in sorted(TUR_ESLEME.items(), key=lambda x: -len(x[0])) if k in kt), ("diger", 0))
        tutar = sayi(_al(r, i["tutar"]))
        if i["yon"] is not None and _al(r, i["yon"]):
            y = katla(_al(r, i["yon"]))
            yon = 1 if y in ("a", "alacak", "giris", "+", "gelen") else -1 if y in ("b", "borc", "cikis", "-", "giden") else yon
        if tutar < 0:
            yon, tutar = -1, -tutar
        if yon == 0:
            uyarilar.append(f"Satır {n}: işlem tipi '{tip}' yönü belirlenemedi (giriş/çıkış); Yön sütunu ekleyin. Atlandı.")
            continue
        pb = (str(_al(r, i["pb"]) or "").strip().upper() or "TL").replace("TRY", "TL")
        if pb != "TL":
            if pb not in kurlar:
                kursuz.add(pb)
                continue
            tutar *= kurlar[pb]
        islemler.append(Islem(str(_al(r, i["no"]) or f"S{n}"), t, musteri, tip, kanal, yon, tutar,
                              str(_al(r, i["karsi"]) or "").strip(), str(_al(r, i["ulke"]) or "").strip().upper(),
                              str(_al(r, i["aciklama"]) or "").strip()))
    if kursuz:
        uyarilar.append(f"Kuru verilmeyen para birimleri ({', '.join(sorted(kursuz))}) atlandı: --kur PB=değer verin.")
    islemler.sort(key=lambda x: (x.musteri, x.zaman, x.no))
    return islemler, uyarilar


def musterileri_oku(yol: Path | None) -> dict[str, dict]:
    if not yol:
        return {}
    s = tablo_oku(yol)
    b = s[0]
    i_no = _bul(b, "müşteri no", "müşteri numarası", "müşteri", "hesap no")
    if i_no is None:
        raise SystemExit(f"Müşteri dosyasında 'Müşteri No' sütunu bulunamadı. Başlıklar: {b}")
    i_ad, i_tip = _bul(b, "ad soyad", "unvan", "ad soyad unvan", "müşteri adı", "ad"), _bul(b, "müşteri tipi", "tip", "tür")
    i_gelir = _bul(b, "beyan aylık gelir", "aylık gelir", "beyan edilen gelir", "aylık ciro")
    i_meslek, i_risk = _bul(b, "meslek", "faaliyet", "sektör"), _bul(b, "risk profili", "risk", "risk seviyesi")
    return {str(_al(r, i_no)).strip(): {"ad": str(_al(r, i_ad) or "").strip(), "tip": str(_al(r, i_tip) or "").strip(),
                                         "gelir": sayi(_al(r, i_gelir)) if _al(r, i_gelir) not in (None, "") else None,
                                         "meslek": str(_al(r, i_meslek) or "").strip(), "risk": str(_al(r, i_risk) or "").strip()}
            for r in s[1:] if _al(r, i_no)}


def ulkeleri_oku(yol: Path | None) -> dict[str, str]:
    if not yol:
        return {}
    sonuc = {}
    for r in tablo_oku(yol)[1:]:
        if r and r[0]:
            sonuc[str(r[0]).strip().upper()] = str(r[1]).strip() if len(r) > 1 and r[1] else ""
    return sonuc


# ----------------------------------------------------------------------------
# Senaryolar
# ----------------------------------------------------------------------------

def alarm(liste, kod, ad, musteri, islemler, tutar, aciklama, agirlik):
    liste.append({"senaryo": kod, "ad": ad, "musteri": musteri, "islemler": [x.no for x in islemler],
                  "bas": min(x.zaman for x in islemler), "son": max(x.zaman for x in islemler),
                  "tutar": tutar, "aciklama": aciklama, "agirlik": agirlik})


def _kanallar(p) -> set[str]:
    return set(p.get("kanallar", ["nakit", "transfer", "swift"]))


def s1_parcali(isl: list[Islem], p: dict, a: list) -> None:
    """Eşiğin hemen altındaki işlemler kısa sürede tekrar ediyor ve toplamı eşiği aşıyor."""
    esik, alt = Decimal(str(p["esik_tl"])), Decimal(str(p["esik_tl"])) * Decimal(str(p["alt_oran"]))
    pencere = timedelta(days=p["pencere_gun"])
    for yon in (1, -1):
        aday = [x for x in isl if x.yon == yon and x.kanal in _kanallar(p) and alt <= x.tutar < esik]
        i = 0
        while i < len(aday):
            grup = [x for x in aday[i:] if x.zaman - aday[i].zaman <= pencere]
            toplam = sum((x.tutar for x in grup), SIFIR)
            if len(grup) >= p["en_az_islem"] and toplam >= esik:
                alarm(a, "S1", "Parçalı işlem", aday[i].musteri, grup, toplam,
                      f"{len(grup)} {'giriş' if yon > 0 else 'çıkış'} işlemi, her biri {tl(alt)}–{tl(esik)} TL aralığında, "
                      f"{p['pencere_gun']} gün içinde; toplam {tl(toplam)} TL", p["agirlik"])
                i += len(grup)
            else:
                i += 1


def s2_hizli_gecis(isl: list[Islem], p: dict, a: list) -> None:
    """Büyük giriş(ler)in ardından kısa sürede büyük oranda çıkış."""
    pencere = timedelta(hours=p["pencere_saat"])
    en_az = Decimal(str(p["en_az_giris_tl"]))
    kullanildi = set()
    for x in isl:
        if x.yon < 0 or x.no in kullanildi:
            continue
        girisler = [y for y in isl if y.yon > 0 and x.zaman <= y.zaman <= x.zaman + pencere]
        g_toplam = sum((y.tutar for y in girisler), SIFIR)
        if g_toplam < en_az:
            continue
        cikislar = [y for y in isl if y.yon < 0 and x.zaman <= y.zaman <= girisler[-1].zaman + pencere]
        c_toplam = sum((y.tutar for y in cikislar), SIFIR)
        if c_toplam >= g_toplam * Decimal(str(p["cikis_orani"])):
            kullanildi.update(y.no for y in girisler)
            sure = (max(y.zaman for y in cikislar) - x.zaman).total_seconds() / 3600
            alarm(a, "S2", "Hızlı geçiş", x.musteri, girisler + cikislar, g_toplam,
                  f"{tl(g_toplam)} TL giriş, ilk girişten sonraki {sure:.0f} saatte {tl(c_toplam)} TL çıkış "
                  f"(çıkış/giriş %{float(c_toplam / g_toplam * 100):.0f})", p["agirlik"])


def s3_nakit(isl: list[Islem], p: dict, a: list) -> None:
    giris = [x for x in isl if x.yon > 0]
    nakit = [x for x in giris if x.kanal == "nakit"]
    tg, tn = sum((x.tutar for x in giris), SIFIR), sum((x.tutar for x in nakit), SIFIR)
    if nakit and tn >= Decimal(str(p["en_az_nakit_tl"])) and tg and tn / tg >= Decimal(str(p["nakit_payi"])):
        alarm(a, "S3", "Nakit yoğunluğu", nakit[0].musteri, nakit, tn,
              f"Dönem girişlerinde nakit payı %{float(tn / tg * 100):.0f} ({len(nakit)} işlem, {tl(tn)} TL)", p["agirlik"])


def s4_ulke(isl: list[Islem], p: dict, a: list, ulkeler: dict[str, str]) -> None:
    for x in isl:
        if x.ulke and x.ulke in ulkeler and x.tutar >= Decimal(str(p.get("en_az_tl", 0))):
            alarm(a, "S4", "Riskli ülke", x.musteri, [x], x.tutar,
                  f"{'Gelen' if x.yon > 0 else 'Giden'} {tl(x.tutar)} TL · ülke {x.ulke}"
                  + (f" ({ulkeler[x.ulke]})" if ulkeler[x.ulke] else ""), p["agirlik"])


def s5_cok_gonderen(isl: list[Islem], p: dict, a: list) -> None:
    pencere = timedelta(days=p["pencere_gun"])
    giris = [x for x in isl if x.yon > 0 and x.kanal in ("transfer", "swift") and x.karsi]
    i = 0
    while i < len(giris):
        grup = [x for x in giris[i:] if x.zaman - giris[i].zaman <= pencere]
        kisi = {katla(x.karsi) for x in grup}
        if len(kisi) >= p["en_az_farkli_gonderen"]:
            toplam = sum((x.tutar for x in grup), SIFIR)
            alarm(a, "S5", "Çok sayıda gönderen", giris[i].musteri, grup, toplam,
                  f"{p['pencere_gun']} gün içinde {len(kisi)} farklı kişiden {len(grup)} giriş, toplam {tl(toplam)} TL", p["agirlik"])
            i += len(grup)
        else:
            i += 1


def s6_profil(isl: list[Islem], p: dict, a: list, musteri: dict | None) -> None:
    gelir = (musteri or {}).get("gelir")
    gecmis: dict[int, list[Decimal]] = {1: [], -1: []}       # giriş ve çıkış geçmişi ayrı tutulur
    for x in isl:
        nedenler = []
        if gelir and x.tutar >= gelir * Decimal(str(p["gelir_kati"])):
            nedenler.append(f"beyan aylık gelirin {float(x.tutar / gelir):.1f} katı".replace(".", ","))
        if len(gecmis[x.yon]) >= p["en_az_gecmis_islem"]:
            ort = median(gecmis[x.yon])
            if ort and x.tutar >= ort * Decimal(str(p["gecmis_kati"])):
                nedenler.append(f"müşterinin önceki {'giriş' if x.yon > 0 else 'çıkış'} medyanının ({tl(ort)} TL) "
                            f"{float(x.tutar / ort):.0f} katı")
        if nedenler and x.tutar >= Decimal(str(p["en_az_tl"])):
            alarm(a, "S6", "Profil dışı tutar", x.musteri, [x], x.tutar, f"{tl(x.tutar)} TL: " + "; ".join(nedenler), p["agirlik"])
        gecmis[x.yon].append(x.tutar)


def s7_uyuyan(isl: list[Islem], p: dict, a: list) -> None:
    for onceki, x in zip(isl, isl[1:]):
        if (x.zaman - onceki.zaman).days >= p["hareketsiz_gun"]:
            sonra = [y for y in isl if x.zaman <= y.zaman <= x.zaman + timedelta(days=p["pencere_gun"])]
            toplam = sum((y.tutar for y in sonra), SIFIR)
            if toplam >= Decimal(str(p["en_az_tl"])):
                alarm(a, "S7", "Uyuyan hesap", x.musteri, sonra, toplam,
                      f"{(x.zaman - onceki.zaman).days} gün hareketsizlik sonrası {p['pencere_gun']} günde {len(sonra)} işlem, "
                      f"{tl(toplam)} TL", p["agirlik"])


def s8_kelime(isl: list[Islem], p: dict, a: list) -> None:
    kelimeler = [katla(k) for k in p.get("kelimeler", []) if k.strip()]
    for x in isl:
        metin = f" {katla(x.aciklama)} {katla(x.karsi)} "
        bulunan = [k for k in kelimeler if f" {k} " in metin or (len(k) >= 5 and k in metin)]
        if bulunan:
            alarm(a, "S8", "Açıklamada anahtar kelime", x.musteri, [x], x.tutar,
                  f"'{', '.join(bulunan)}' · {x.aciklama[:80]}", p["agirlik"])


def s9_yuvarlak(isl: list[Islem], p: dict, a: list) -> None:
    taban = Decimal(str(p["yuvarlak_taban_tl"]))
    pencere = timedelta(days=p["pencere_gun"])
    aday = [x for x in isl if x.tutar >= taban and x.tutar % taban == 0]
    i = 0
    while i < len(aday):
        grup = [x for x in aday[i:] if x.zaman - aday[i].zaman <= pencere]
        if len(grup) >= p["en_az_islem"]:
            toplam = sum((x.tutar for x in grup), SIFIR)
            alarm(a, "S9", "Yuvarlak tutar tekrarı", aday[i].musteri, grup, toplam,
                  f"{p['pencere_gun']} günde {len(grup)} adet {tl(taban)} TL'nin katı işlem, toplam {tl(toplam)} TL", p["agirlik"])
            i += len(grup)
        else:
            i += 1


def calistir(islem_yolu: Path, senaryo_yolu: Path, cikti: Path, musteri_yolu: Path | None = None,
             ulke_yolu: Path | None = None, kurlar: dict[str, Decimal] | None = None) -> dict:
    ayar = json.loads(senaryo_yolu.read_text(encoding="utf-8"))
    sen = {k: v for k, v in ayar["senaryolar"].items() if v.get("aktif", True)}
    islemler, uyarilar = islemleri_oku(islem_yolu, kurlar or {})
    musteriler, ulkeler = musterileri_oku(musteri_yolu), ulkeleri_oku(ulke_yolu)
    if "S4" in sen and not ulkeler:
        uyarilar.append("S4 (riskli ülke) için ülke listesi verilmedi; senaryo çalışmadı.")
    if "S6" in sen and not musteriler:
        uyarilar.append("Müşteri dosyası yok: S6 yalnız müşterinin kendi işlem geçmişine göre çalıştı.")
    gruplar: dict[str, list[Islem]] = defaultdict(list)
    for x in islemler:
        gruplar[x.musteri].append(x)
    alarmlar: list[dict] = []
    for m, isl in gruplar.items():
        if "S1" in sen:
            s1_parcali(isl, sen["S1"], alarmlar)
        if "S2" in sen:
            s2_hizli_gecis(isl, sen["S2"], alarmlar)
        if "S3" in sen:
            s3_nakit(isl, sen["S3"], alarmlar)
        if "S4" in sen and ulkeler:
            s4_ulke(isl, sen["S4"], alarmlar, ulkeler)
        if "S5" in sen:
            s5_cok_gonderen(isl, sen["S5"], alarmlar)
        if "S6" in sen:
            s6_profil(isl, sen["S6"], alarmlar, musteriler.get(m))
        if "S7" in sen:
            s7_uyuyan(isl, sen["S7"], alarmlar)
        if "S8" in sen:
            s8_kelime(isl, sen["S8"], alarmlar)
        if "S9" in sen:
            s9_yuvarlak(isl, sen["S9"], alarmlar)
    for n, x in enumerate(sorted(alarmlar, key=lambda x: (x["musteri"], x["bas"], x["senaryo"])), 1):
        x["no"] = f"A{n:04d}"
    ozet = musteri_ozeti(alarmlar, gruplar, musteriler, ayar.get("risk_puani", {}))
    _rapor(alarmlar, ozet, islemler, sen, ayar, uyarilar, cikti)
    return {"alarmlar": alarmlar, "musteriler": ozet, "islemler": islemler, "uyarilar": uyarilar}


def musteri_ozeti(alarmlar, gruplar, musteriler, risk_ayar) -> list[dict]:
    yuksek, orta = risk_ayar.get("yuksek", 60), risk_ayar.get("orta", 30)
    m: dict[str, dict] = {}
    for x in alarmlar:
        k = m.setdefault(x["musteri"], {"musteri": x["musteri"], "senaryolar": set(), "alarm": 0, "puan": 0})
        k["alarm"] += 1
        if x["senaryo"] not in k["senaryolar"]:          # aynı senaryo bir kez puanlanır; farklı senaryolar toplanır
            k["puan"] += x["agirlik"]
        k["senaryolar"].add(x["senaryo"])
    for k in m.values():
        k["puan"] = min(100, k["puan"])
        k["oncelik"] = "Yüksek" if k["puan"] >= yuksek else "Orta" if k["puan"] >= orta else "Düşük"
        isl = gruplar[k["musteri"]]
        k["giris"] = sum((x.tutar for x in isl if x.yon > 0), SIFIR)
        k["cikis"] = sum((x.tutar for x in isl if x.yon < 0), SIFIR)
        k["islem"] = len(isl)
        k.update({f"m_{a}": v for a, v in (musteriler.get(k["musteri"]) or {}).items()})
    return sorted(m.values(), key=lambda k: (-k["puan"], k["musteri"]))


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
ONCELIK = {"Yüksek": "F8C9C6", "Orta": "FFF4CE", "Düşük": "E8F0FE"}
PARA = "#,##0.00"
UST = Alignment(vertical="top", wrap_text=True)
UYARI = ("Bu liste inceleme önerisidir; bir işlemin şüpheli olup olmadığına uyum biriminin değerlendirmesiyle karar "
         "verilir. Şüpheli işlem bildirimi yapıldığı veya yapılacağı bilgisi, işleme taraf olanlar dahil kimseye "
         "açıklanamaz (5549 sayılı Kanun md. 4/2). Senaryo eşikleri örnektir; kurumunuzun risk politikası ve güncel "
         "MASAK düzenlemelerine göre belirleyin.")


def _baslik(ws, satir=1):
    for c in ws[satir]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI


def _rapor(alarmlar, ozet, islemler, sen, ayar, uyarilar, cikti):
    wb = Workbook()
    o = wb.active
    o.title = "Özet"
    o.append(["Şüpheli işlem senaryo taraması — inceleme listesi"])
    o["A1"].font = Font(bold=True, size=12)
    if islemler:
        bas, son = min(x.zaman for x in islemler), max(x.zaman for x in islemler)
        o.append(["Dönem", f"{bas:%d.%m.%Y} – {son:%d.%m.%Y}"])
    o.append(["Taranan işlem", len(islemler)])
    o.append(["Taranan müşteri", len({x.musteri for x in islemler})])
    o.append(["Alarm", len(alarmlar)])
    o.append(["Alarmlı müşteri", len(ozet)])
    for p in ("Yüksek", "Orta", "Düşük"):
        o.append([f"Öncelik: {p}", sum(1 for k in ozet if k["oncelik"] == p)])
        o.cell(o.max_row, 1).fill = PatternFill("solid", fgColor=ONCELIK[p])
    o.append([])
    o.append(["Senaryo", "Alarm", "Parametreler"])
    _baslik(o, o.max_row)
    for kod, p in sen.items():
        o.append([f"{kod} {p.get('ad', '')}", sum(1 for x in alarmlar if x["senaryo"] == kod),
                  ", ".join(f"{k}={v}" for k, v in p.items() if k not in ("ad", "aktif", "aciklama", "kelimeler"))
                  + (f", {len(p['kelimeler'])} kelime" if "kelimeler" in p else "")])
    for u in uyarilar:
        o.append(["Uyarı", u])
    o.append([])
    o.append(["ÖNEMLİ", UYARI])
    o.cell(o.max_row, 1).font = Font(bold=True, color="B00020")
    o.column_dimensions["A"].width = 30
    o.column_dimensions["B"].width = 12
    o.column_dimensions["C"].width = 110
    o.cell(o.max_row, 2).alignment = UST
    o.merge_cells(start_row=o.max_row, start_column=2, end_row=o.max_row, end_column=3)

    m = wb.create_sheet("Müşteri Öncelik")
    m.append(["Müşteri No", "Ad / Unvan", "Tip", "Meslek / Faaliyet", "Beyan Aylık Gelir", "Mevcut Risk Profili", "Puan", "Öncelik",
              "Senaryolar", "Alarm", "İşlem", "Toplam Giriş", "Toplam Çıkış", "İnceleyen", "Karar (Kapat / İzle / Bildirim)",
              "Gerekçe"])
    _baslik(m)
    for k in ozet:
        m.append([k["musteri"], k.get("m_ad", ""), k.get("m_tip", ""), k.get("m_meslek", ""),
                  float(k["m_gelir"]) if k.get("m_gelir") is not None else None, k.get("m_risk", ""), k["puan"], k["oncelik"],
                  ", ".join(sorted(k["senaryolar"])), k["alarm"], k["islem"], float(k["giris"]), float(k["cikis"]), None, None, None])
        m.cell(m.max_row, 8).fill = PatternFill("solid", fgColor=ONCELIK[k["oncelik"]])
        for c in (5, 12, 13):
            m.cell(m.max_row, c).number_format = PARA
    for j, w in enumerate((12, 24, 10, 18, 14, 12, 7, 9, 16, 7, 7, 15, 15, 14, 22, 40), 1):
        m.column_dimensions[get_column_letter(j)].width = w
    m.freeze_panes = "C2"
    m.auto_filter.ref = m.dimensions

    a = wb.create_sheet("Alarmlar")
    a.append(["Alarm No", "Müşteri No", "Senaryo", "Senaryo Adı", "Başlangıç", "Bitiş", "Tutar (TL)", "Açıklama", "İşlem Sayısı",
              "İşlemler", "Değerlendirme"])
    _baslik(a)
    for x in sorted(alarmlar, key=lambda x: x["no"]):
        a.append([x["no"], x["musteri"], x["senaryo"], x["ad"], x["bas"], x["son"], float(x["tutar"]), x["aciklama"],
                  len(x["islemler"]), ", ".join(x["islemler"][:30]) + (" …" if len(x["islemler"]) > 30 else ""), None])
        for c in (5, 6):
            a.cell(a.max_row, c).number_format = "DD.MM.YYYY HH:MM"
        a.cell(a.max_row, 7).number_format = PARA
        a.cell(a.max_row, 8).alignment = UST
    for j, w in enumerate((9, 12, 8, 22, 16, 16, 15, 70, 8, 40, 30), 1):
        a.column_dimensions[get_column_letter(j)].width = w
    a.freeze_panes = "C2"
    a.auto_filter.ref = a.dimensions

    d = wb.create_sheet("Alarmlı İşlemler")
    d.append(["Alarm No", "Senaryo", "Müşteri No", "İşlem No", "Tarih", "İşlem Tipi", "Yön", "Tutar (TL)", "Karşı Taraf", "Ülke", "Açıklama"])
    _baslik(d)
    isl = {x.no: x for x in islemler}
    for x in sorted(alarmlar, key=lambda x: x["no"]):
        for no in x["islemler"]:
            y = isl[no]
            d.append([x["no"], x["senaryo"], y.musteri, y.no, y.zaman, y.tip, "Giriş" if y.yon > 0 else "Çıkış", float(y.tutar),
                      y.karsi, y.ulke, y.aciklama])
            d.cell(d.max_row, 5).number_format = "DD.MM.YYYY HH:MM"
            d.cell(d.max_row, 8).number_format = PARA
    for j, w in enumerate((9, 8, 12, 12, 16, 18, 7, 14, 26, 6, 40), 1):
        d.column_dimensions[get_column_letter(j)].width = w
    d.freeze_panes = "E2"
    d.auto_filter.ref = d.dimensions

    b = wb.create_sheet("Senaryo Tanımları")
    b.append(["Kod", "Senaryo", "Açıklama", "Ağırlık"])
    _baslik(b)
    for kod, p in ayar["senaryolar"].items():
        b.append([kod, p.get("ad", ""), p.get("aciklama", ""), p.get("agirlik", 0)])
        b.cell(b.max_row, 3).alignment = UST
    b.append([])
    b.append(["Puan", "Müşteri puanı: alarm veren her farklı senaryonun ağırlığı toplanır (en çok 100); "
                      f"≥ {ayar.get('risk_puani', {}).get('yuksek', 60)} yüksek, ≥ {ayar.get('risk_puani', {}).get('orta', 30)} orta öncelik"])
    b.column_dimensions["A"].width = 8
    b.column_dimensions["B"].width = 26
    b.column_dimensions["C"].width = 110
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def kur_coz(degerler: list[str] | None) -> dict[str, Decimal]:
    kurlar = {}
    for d in degerler or []:
        m = re.match(r"^([A-Za-z]{3})=([\d.,]+)$", d.strip())
        if not m:
            raise SystemExit(f"Kur biçimi PB=değer olmalı (ör. USD=41,20): {d}")
        kurlar[m.group(1).upper()] = sayi(m.group(2))
    return kurlar


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ornek = BURASI / "ornek_veri"
    ap = argparse.ArgumentParser(description="İşlem verisini yapılandırılabilir senaryolarla tarayıp inceleme listesi çıkarır.")
    ap.add_argument("--islemler", type=Path, default=ornek / "islemler.csv",
                    help="İşlemler (.xlsx/.csv): İşlem No, Tarih, Müşteri No, İşlem Tipi, Tutar [, Para Birimi, Yön, Karşı Taraf, Karşı Taraf Ülke, Açıklama]")
    ap.add_argument("--senaryolar", type=Path, default=BURASI / "senaryolar.json", help="Senaryo eşikleri ve ağırlıkları (JSON)")
    ap.add_argument("--musteriler", type=Path, help="Müşteriler: Müşteri No, Ad/Unvan, Tip, Meslek, Beyan Aylık Gelir, Risk Profili")
    ap.add_argument("--riskli-ulkeler", type=Path, help="Riskli ülke listesi (.csv/.xlsx: Ülke Kodu, Açıklama)")
    ap.add_argument("--kur", nargs="*", help="Döviz kurları, ör. --kur USD=41,20 EUR=48,10")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "senaryo_taramasi.xlsx")
    a = ap.parse_args(argv)
    if a.islemler == ornek / "islemler.csv":
        a.musteriler = a.musteriler or ornek / "musteriler.csv"
        a.riskli_ulkeler = a.riskli_ulkeler or ornek / "riskli_ulkeler.csv"
        a.kur = a.kur or ["USD=41,20", "EUR=48,10"]
    s = calistir(a.islemler, a.senaryolar, a.cikti, a.musteriler, a.riskli_ulkeler, kur_coz(a.kur))
    from collections import Counter
    sayac = Counter(x["senaryo"] for x in s["alarmlar"])
    print(f"[OK] {len(s['islemler'])} işlem · {len(s['alarmlar'])} alarm (" + ", ".join(f"{k}: {n}" for k, n in sorted(sayac.items()))
          + f") · {len(s['musteriler'])} müşteri incelemeye")
    for k in s["musteriler"][:5]:
        print(f"    {k['musteri']}: puan {k['puan']} ({k['oncelik']}) · {', '.join(sorted(k['senaryolar']))}")
    for u in s["uyarilar"]:
        print(f"[!] {u}")
    print("[i] İnceleme önerisidir; bildirim yapıldığı bilgisi kimseye açıklanamaz (5549 s. K. md. 4/2).")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
