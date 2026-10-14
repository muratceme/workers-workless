"""
Öğrenim Ücreti ve İndirim Hesabı — Workers / Workless kod bloğu
Kayıt Kabul ve Öğrenci İşleri › Öğrenci İşleri Sorumlusu
Okulun ücret ve indirim kurallarına göre öğrenci bazında net öğrenim ücretini ve taksit planını hesaplar.
  Brüt ücret   = ücret tablosunda öğrencinin kademe / sınıfına ait yıllık ücret
  İndirimler   = kardeş (2. çocuk / 3. ve sonraki), erken kayıt, peşin ödeme (tek taksit), personel çocuğu, burs %, diğer %
  Birleştirme  = zincir (varsayılan): her indirim bir öncekinden kalan tutara uygulanır; toplam: oranlar toplanır
  Tavan        = toplam indirim oranı --azami (%50) ile sınırlanır (burs dahil; --burs-haric ile burs tavana girmez)
  Taksit       = peşinat --pesinat % + kalan tutar eşit aylık taksitler (son taksit kuruş farkını kapatır)
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
UCRET = {"kademe": ("kademe", "okul", "seviye"), "sinif": ("sinif", "sinif seviyesi"), "ucret": ("yillik ucret", "ucret", "egitim ucreti", "ogrenim ucreti")}
INDIRIM = {"kod": ("kod", "indirim kodu"), "ad": ("indirim", "ad", "aciklama"), "oran": ("oran", "oran yuzde", "indirim orani")}
OGRENCI = {"no": ("ogrenci no", "no", "kayit no"), "ad": ("ad soyad", "ogrenci", "ogrenci adi"), "kademe": ("kademe", "okul", "seviye"), "sinif": ("sinif",),
           "kardes": ("kardes sirasi", "kardes", "kacinci cocuk"), "erken": ("erken kayit",), "personel": ("personel cocugu", "personel"), "burs": ("burs", "burs orani"),
           "diger": ("diger indirim", "ozel indirim"), "taksit": ("taksit sayisi", "taksit")}
KODLAR = ("KARDES2", "KARDES3", "ERKEN", "PESIN", "PERSONEL")
SIFIR = Decimal(0)


def ucret_oku(yol: Path) -> dict:
    d = {}
    for r in kayitlar(yol, UCRET, ("kademe", "ucret")):
        if not metin(r["kademe"]):
            continue
        u = para(r["ucret"])
        if u is None or u <= 0:
            raise ValueError(f"{yol.name} satır {r['_satir']}: ücret okunamadı.")
        d[(katla(r["kademe"]), katla(r["sinif"]) if katla(r["sinif"]) not in ("", "tumu", "hepsi") else "")] = u
    if not d:
        raise ValueError(f"{yol.name}: ücret satırı yok.")
    return d


def indirim_oku(yol: Path) -> dict:
    d = {}
    for r in kayitlar(yol, INDIRIM, ("kod", "oran")):
        k = metin(r["kod"]).upper()
        if not k:
            continue
        o = para(r["oran"])
        if k not in KODLAR or o is None or not 0 <= o <= 100:
            raise ValueError(f"{yol.name} satır {r['_satir']}: kod ({', '.join(KODLAR)}) ya da oran okunamadı.")
        d[k] = {"ad": metin(r["ad"]) or k, "oran": o}
    return d


def ogrenci_oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, OGRENCI, ("no", "kademe")):
        no = metin(r["no"])
        if not no:
            continue
        kardes = para(r["kardes"]) if metin(r["kardes"]) else Decimal(1)
        burs, diger = (para(r[a]) if metin(r[a]) else SIFIR for a in ("burs", "diger"))
        taksit = para(r["taksit"]) if metin(r["taksit"]) else None
        if kardes is None or kardes < 1 or burs is None or not 0 <= burs <= 100 or diger is None or not 0 <= diger <= 100 or (metin(r["taksit"]) and (taksit is None or taksit < 1 or taksit != int(taksit))):
            uy.append(uyari("Yüksek", "Satır okunamadı", f"{no} {metin(r['ad'])}".strip(), f"Satır {r['_satir']}: kardeş sırası, burs, diğer indirim ya da taksit sayısı okunamadı."))
            continue
        e = lambda a: katla(r[a]) in ("evet", "e", "x", "var", "1")
        l.append({"no": no, "ad": metin(r["ad"]), "kademe": metin(r["kademe"]), "sinif": metin(r["sinif"]), "kardes": int(kardes), "erken": e("erken"), "personel": e("personel"),
                  "burs": burs, "diger": diger, "taksit": int(taksit) if taksit else None})
    return l, uy


def ay_ekle(g: date, ay: int) -> date:
    y, m = divmod(g.month - 1 + ay, 12)
    yil, ayn = g.year + y, m + 1
    son = (date(yil + (ayn == 12), ayn % 12 + 1, 1) - timedelta(days=1)).day
    return date(yil, ayn, min(g.day, son))


def hesapla(ogrenciler: list[dict], ucretler: dict, indirimler: dict, ilk: date, varsayilan_taksit: int = 8, pesinat: Decimal = SIFIR, birlesim: str = "zincir",
            azami: Decimal = Decimal(50), burs_haric: bool = False) -> dict:
    if not ogrenciler:
        raise ValueError("Okunabilen öğrenci yok.")
    uy, sonuc = [], []
    for o in ogrenciler:
        kim = f"{o['no']} {o['ad']}".strip()
        brut = ucretler.get((katla(o["kademe"]), katla(o["sinif"]))) or ucretler.get((katla(o["kademe"]), ""))
        if brut is None:
            uy.append(uyari("Yüksek", "Ücret bulunamadı", kim, f"'{o['kademe']} {o['sinif']}' için ücret tablosunda satır yok."))
            continue
        taksit = o["taksit"] or varsayilan_taksit
        uyg = []
        if o["kardes"] == 2 and "KARDES2" in indirimler:
            uyg.append(("KARDES2", indirimler["KARDES2"]))
        elif o["kardes"] >= 3 and "KARDES3" in indirimler:
            uyg.append(("KARDES3", indirimler["KARDES3"]))
        if o["erken"] and "ERKEN" in indirimler:
            uyg.append(("ERKEN", indirimler["ERKEN"]))
        if taksit == 1 and "PESIN" in indirimler:
            uyg.append(("PESIN", indirimler["PESIN"]))
        if o["personel"] and "PERSONEL" in indirimler:
            uyg.append(("PERSONEL", indirimler["PERSONEL"]))
        if o["diger"]:
            uyg.append(("DIGER", {"ad": "Diğer indirim", "oran": o["diger"]}))
        burs = [("BURS", {"ad": "Burs", "oran": o["burs"]})] if o["burs"] else []
        tavana = uyg + ([] if burs_haric else burs)
        if birlesim == "toplam":
            oran_t = sum((x["oran"] for _, x in tavana), SIFIR)
        else:
            kalan = Decimal(1)
            for _, x in tavana:
                kalan *= 1 - x["oran"] / 100
            oran_t = (1 - kalan) * 100
        sinirli = oran_t > azami
        oran_t = min(oran_t, azami)
        net = yuvarla(brut * (1 - oran_t / 100))
        if burs_haric and burs:
            net = yuvarla(net * (1 - o["burs"] / 100))
        if sinirli:
            uy.append(uyari("Bilgi", "İndirim tavana takıldı", kim, f"İndirimler toplamı %{azami} tavanı ile sınırlandı ({', '.join(x['ad'] for _, x in tavana)})."))
        if o["burs"] == 100:
            uy.append(uyari("Bilgi", "Tam burslu", kim, "Burs %100" + ("; tavan nedeniyle ücretin bir kısmı kalıyor, burs koşullarını kontrol edin." if net > 0 else ".")))
        ps = yuvarla(net * pesinat / 100) if taksit > 1 else SIFIR
        kalan_t = net - ps
        tk = yuvarla(kalan_t / taksit) if taksit else SIFIR
        plan = ([{"no": "Peşinat", "vade": ilk, "tutar": ps}] if ps else [])
        for i in range(taksit):
            plan.append({"no": i + 1, "vade": ay_ekle(ilk, i + (1 if ps else 0)), "tutar": tk if i < taksit - 1 else kalan_t - tk * (taksit - 1)})
        sonuc.append({**o, "brut": brut, "indirimler": [x["ad"] for _, x in uyg + burs], "oran": yuvarla(Decimal(100) - net * 100 / brut), "indirim": brut - net, "net": net,
                      "taksit_sayisi": taksit, "plan": plan, "sinirli": sinirli})
    ozet = {}
    for s in sonuc:
        z = ozet.setdefault(s["kademe"], {"kademe": s["kademe"], "ogrenci": 0, "brut": SIFIR, "indirim": SIFIR, "net": SIFIR})
        z["ogrenci"] += 1
        z["brut"] += s["brut"]
        z["indirim"] += s["indirim"]
        z["net"] += s["net"]
    for z in ozet.values():
        z["oran"] = yuvarla(z["indirim"] * 100 / z["brut"]) if z["brut"] else None
    return {"ogrenciler": sonuc, "ozet": list(ozet.values()), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Öğrenci Ücretleri"
    _baslik(ws, ["Öğrenci No", "Ad Soyad", "Kademe", "Sınıf", "Brüt Ücret", "Uygulanan İndirimler", "Toplam İndirim %", "İndirim Tutarı", "Net Ücret", "Taksit", "Taksit Tutarı"],
            (10, 22, 10, 7, 13, 44, 9, 13, 13, 7, 12), "C2")
    for s in r["ogrenciler"]:
        _satir(ws, [s["no"], s["ad"], s["kademe"], s["sinif"], s["brut"], ", ".join(s["indirimler"]) or "-", s["oran"], s["indirim"], s["net"], s["taksit_sayisi"],
                    next((p["tutar"] for p in s["plan"] if p["no"] == 1), None)], para_s=(5, 7, 8, 9, 11), sar=(6,), dolgu={7: SARI if s["sinirli"] else None})
    ws = wb.create_sheet("Taksit Planı")
    _baslik(ws, ["Öğrenci No", "Ad Soyad", "Taksit", "Vade", "Tutar"], (10, 22, 9, 11, 13))
    for s in r["ogrenciler"]:
        for p in s["plan"]:
            _satir(ws, [s["no"], s["ad"], p["no"], p["vade"], p["tutar"]], para_s=(5,), tarih_s=(4,))
    ws = wb.create_sheet("Kademe Özeti")
    _baslik(ws, ["Kademe", "Öğrenci", "Brüt", "İndirim", "Net", "İndirim Payı %"], (12, 8, 15, 14, 15, 10))
    for z in r["ozet"]:
        _satir(ws, [z["kademe"], z["ogrenci"], z["brut"], z["indirim"], z["net"], z["oran"]], para_s=(3, 4, 5, 6))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(ucret_yol: Path, indirim_yol: Path, ogrenci_yol: Path, cikti: Path, ilk: date, **k) -> dict:
    u, i = ucret_oku(ucret_yol), indirim_oku(indirim_yol)
    o, uy = ogrenci_oku(ogrenci_yol)
    r = hesapla(o, u, i, ilk, **k)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Öğrenci bazında indirimli öğrenim ücretini ve taksit planını hesaplar.")
    p.add_argument("--ucretler", type=Path, help="Kademe, Sınıf (boş = tümü), Yıllık Ücret")
    p.add_argument("--indirimler", type=Path, help="Kod (KARDES2 / KARDES3 / ERKEN / PESIN / PERSONEL), İndirim, Oran %%")
    p.add_argument("--ogrenciler", type=Path, help="Öğrenci No, Ad Soyad, Kademe, Sınıf, Kardeş Sırası, Erken Kayıt, Personel Çocuğu, Burs %%, Diğer İndirim %%, Taksit Sayısı")
    p.add_argument("--ilk-taksit", help="İlk taksit (ya da peşinat) tarihi GG.AA.YYYY (örnek veride 01.09.2026)")
    p.add_argument("--taksit", type=int, default=8, help="Satırda yazmıyorsa taksit sayısı (varsayılan 8)")
    p.add_argument("--pesinat", default="0", help="Taksitli ödemede peşinat yüzdesi (varsayılan 0)")
    p.add_argument("--birlesim", choices=("zincir", "toplam"), default="zincir", help="İndirimler zincirleme mi uygulanır, toplanır mı (varsayılan zincir)")
    p.add_argument("--azami", default="50", help="Toplam indirim tavanı yüzdesi (varsayılan 50)")
    p.add_argument("--burs-haric", action="store_true", help="Burs tavana girmez; diğer indirimlerden sonra ayrıca uygulanır")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "ogrenim_ucretleri.xlsx")
    a = p.parse_args(argv)
    ornek = a.ogrenciler is None
    yollar = [a.ucretler or ORNEK / "ucretler.csv", a.indirimler or ORNEK / "indirimler.csv", a.ogrenciler or ORNEK / "ogrenciler.csv"]
    ilk = tarih(a.ilk_taksit) if a.ilk_taksit else (date(2026, 9, 1) if ornek else None)
    ps, az = para(a.pesinat), para(a.azami)
    if ilk is None or not 1 <= a.taksit <= 24 or ps is None or not 0 <= ps < 100 or az is None or not 0 <= az <= 100:
        print("[X] --ilk-taksit GG.AA.YYYY zorunlu (örnek veri dışında); --taksit 1-24, --pesinat 0-99, --azami 0-100 olmalı")
        return 2
    for y in yollar:
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(*yollar, a.cikti, ilk, varsayilan_taksit=a.taksit, pesinat=ps, birlesim=a.birlesim, azami=az, burs_haric=a.burs_haric)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    for z in r["ozet"]:
        print(f"[OK] {z['kademe']}: {z['ogrenci']} öğrenci · brüt {z['brut']:,.2f} · indirim {z['indirim']:,.2f} (%{z['oran']}) · net {z['net']:,.2f}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
