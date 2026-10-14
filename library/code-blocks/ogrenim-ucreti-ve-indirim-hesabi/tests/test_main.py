import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
UCRET = {("lise", ""): D(600000), ("lise", "hazirlik"): D(500000)}
IND = {"KARDES2": {"ad": "Kardeş 2", "oran": D(10)}, "KARDES3": {"ad": "Kardeş 3", "oran": D(15)}, "ERKEN": {"ad": "Erken", "oran": D(5)}, "PESIN": {"ad": "Peşin", "oran": D(7)},
       "PERSONEL": {"ad": "Personel", "oran": D(50)}}
ILK = date(2026, 9, 1)


def o(no, sinif="10", kardes=1, erken=False, personel=False, burs=0, diger=0, taksit=None):
    return {"no": no, "ad": "", "kademe": "Lise", "sinif": sinif, "kardes": kardes, "erken": erken, "personel": personel, "burs": D(burs), "diger": D(diger), "taksit": taksit}


class HesapTesti(unittest.TestCase):
    def test_indirimler(self):
        l = [o("1"), o("2", kardes=2, erken=True), o("3", "Hazırlık", kardes=3, taksit=1), o("4", personel=True, burs=50), o("5", diger=5, taksit=3)]
        r = main.hesapla(l, UCRET, IND, ILK, 8, D(20))
        s = {x["no"]: x for x in r["ogrenciler"]}
        # 2: 600.000 × 0,90 × 0,95 = 513.000 (zincir)
        self.assertEqual([s[n]["net"] for n in "12345"], [D(600000), D(513000), D("395250.00"), D(300000), D(570000)])
        self.assertEqual((s["3"]["indirimler"], s["4"]["oran"], s["4"]["sinirli"]), (["Kardeş 3", "Peşin"], D(50), True))
        # taksit: 2 → %20 peşinat 102.600 + 8 × 51.300
        p = s["2"]["plan"]
        self.assertEqual((p[0]["no"], p[0]["tutar"], p[1]["tutar"], p[1]["vade"], p[-1]["vade"], len(p)), ("Peşinat", D(102600), D(51300), date(2026, 10, 1), date(2027, 5, 1), 9))
        self.assertEqual([x["tutar"] for x in s["3"]["plan"]], [D("395250.00")])
        self.assertEqual(sum(x["tutar"] for x in s["5"]["plan"]), D(570000))
        self.assertEqual([u["tur"] for u in r["uyarilar"]], ["İndirim tavana takıldı"])
        self.assertEqual(main.hesapla([o("2", kardes=2, erken=True)], UCRET, IND, ILK, birlesim="toplam")[ "ogrenciler"][0]["net"], D(510000))
        r = main.hesapla([o("4", personel=True, burs=50)], UCRET, IND, ILK, burs_haric=True)
        self.assertEqual((r["ogrenciler"][0]["net"], r["uyarilar"]), (D(150000), []))
        r = main.hesapla([o("6", burs=100), o("7")], UCRET | {("lise", "10"): D(1)}, IND, ILK, azami=D(100))
        self.assertEqual((r["ogrenciler"][0]["net"], r["ozet"][0]["ogrenci"]), (D(0), 2))
        with self.assertRaises(ValueError):
            main.hesapla([], UCRET, IND, ILK)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "ucretler.csv", ORNEK / "indirimler.csv", ORNEK / "ogrenciler.csv", cikti, ILK)
            self.assertEqual([(z["kademe"], z["ogrenci"], z["net"]) for z in r["ozet"]],
                             [("Anaokulu", 7, D("1935675.00")), ("Lise", 9, D("4625200.00")), ("Ortaokul", 15, D("6955122.00")), ("İlkokul", 12, D("4927260.00"))])
            self.assertEqual([u["tur"] for u in r["uyarilar"] if u["onem"] == "Yüksek"], ["Satır okunamadı", "Ücret bulunamadı"])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Öğrenci Ücretleri", "Taksit Planı", "Kademe Özeti", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            u, i, g = Path(tmp) / "u.csv", Path(tmp) / "i.csv", Path(tmp) / "g.csv"
            u.write_text("Okul;Sınıf;Ücret\nLise;Tümü;100.000\n", encoding="utf-8")
            i.write_text("Kod;Oran\nkardes2;10\n", encoding="utf-8")
            g.write_text("No;Kademe;Kardeş;Burs;Taksit\n1;Lise;2;;2\n2;Lise;;120;\n3;Lise;;;1,5\n", encoding="utf-8")
            self.assertEqual((main.ucret_oku(u), main.indirim_oku(i)), ({("lise", ""): D(100000)}, {"KARDES2": {"ad": "KARDES2", "oran": D(10)}}))
            l, uy = main.ogrenci_oku(g)
            self.assertEqual(([(x["no"], x["kardes"], x["taksit"]) for x in l], len(uy)), ([("1", 2, 2)], 2))
            for yol, metin in ((u, "Okul;Ücret\nLise;x\n"), (i, "Kod;Oran\nX;5\n"), (u, "Okul;Ücret\n")):
                yol.write_text(metin, encoding="utf-8")
                with self.assertRaises(ValueError):
                    (main.ucret_oku if yol == u else main.indirim_oku)(yol)
            u.write_text("Okul;Sınıf;Ücret\nLise;Tümü;100.000\n", encoding="utf-8")
            i.write_text("Kod;Oran\nKARDES2;10\n", encoding="utf-8")
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--ucretler", str(u), "--indirimler", str(i), "--ogrenciler", str(g), "--cikti", c]
            self.assertEqual(main.main(arg + ["--ilk-taksit", "01.09.2026", "--taksit", "10", "--pesinat", "25", "--birlesim", "toplam", "--azami", "40", "--burs-haric"]), 0)
            self.assertEqual([main.main(arg + v) for v in ([], ["--ilk-taksit", "x"], ["--ilk-taksit", "01.09.2026", "--taksit", "0"], ["--ilk-taksit", "01.09.2026", "--azami", "101"])], [2] * 4)
            self.assertEqual(main.main(["--ogrenciler", str(Path(tmp) / "yok.csv"), "--ilk-taksit", "01.09.2026"]), 1)
            i.write_text("Kod;Oran\nX;1\n", encoding="utf-8")
            self.assertEqual(main.main(arg + ["--ilk-taksit", "01.09.2026"]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
