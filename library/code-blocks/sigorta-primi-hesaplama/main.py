"""
Sigorta Primi Hesaplama — Workers / Workless kod bloğu
Teknik (Risk Kabul / Underwriting) › Teknik Uzman (Underwriter)

Sizin verdiğiniz tarife fiyatları ve indirim-sürprim tablosuyla teminat bazında prim hesaplar.
  Baz prim     = bedel × fiyat (‰) ÷ 1000
  Süre katsayısı = gün ÷ 365   (tam bir yıllık poliçede 1; kısa süre tarifesi uygulanmaz, gün esası)
  Çarpan       = zincirleme: Π(1 + sürprim) × Π(1 − indirim)     |  toplamsal: 1 + Σ sürprim − Σ indirim
  Net prim     = baz prim × süre katsayısı × çarpan;  asgari primin altındaysa asgari prim
  Vergi        = net prim × --vergi %   (varsayılan 0; ör. BSMV için oranı siz verirsiniz)
Tarife fiyatları koda gömülü değildir. İndirim tavanı --azami-indirim ile sınırlandırılabilir.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


TARIFE_SUTUN = {"teminat": ("teminat",), "sinif": ("risk sinifi", "sinif", "tarife sinifi"), "fiyat": ("fiyat", "fiyat binde", "oran", "tarife fiyati", "binde"),
                "asgari": ("asgari prim", "minimum prim")}
INDIRIM_SUTUN = {"kod": ("kod",), "ad": ("ad", "aciklama"), "tur": ("tur", "tip"), "oran": ("oran", "oran yuzde", "yuzde"), "teminat": ("teminat", "uygulanan teminat")}
RISK_SUTUN = {"no": ("teklif no", "talep no", "police no"), "sigortali": ("sigortali", "musteri"), "sinif": ("risk sinifi", "sinif", "tarife sinifi"), "teminat": ("teminat",),
              "bedel": ("bedel", "sigorta bedeli", "teminat bedeli"), "bas": ("baslangic", "baslangic tarihi"), "bit": ("bitis", "bitis tarihi"), "gun": ("sure gun", "sure", "gun"),
              "kodlar": ("indirim surprim", "indirim surprim kodlari", "kodlar", "indirimler")}
Z, BIR = Decimal(0), Decimal(1)


def tarife_oku(yol: Path) -> tuple[dict, list[dict]]:
    tarife, uy = {}, []
    for r in kayitlar(yol, TARIFE_SUTUN, ("teminat", "fiyat")):
        tem = metin(r["teminat"])
        if not tem:
            continue
        f, sinif = para(r["fiyat"]), katla(r["sinif"]) or "*"
        if f is None or f < 0:
            uy.append(uyari("Yüksek", "Tarife fiyatı okunamadı", tem, f"{yol.name} satır {r['_satir']}: '{metin(r['fiyat'])}'. Satır tarifeye alınmadı."))
            continue
        if (katla(tem), sinif) in tarife:
            uy.append(uyari("Yüksek", "Tarife satırı tekrar ediyor", tem, f"{yol.name} satır {r['_satir']}: aynı teminat ve risk sınıfı iki kez tanımlı; ilk satır kullanıldı."))
            continue
        tarife[(katla(tem), sinif)] = {"fiyat": f, "asgari": para(r["asgari"])}
    if not tarife:
        raise ValueError(f"{yol.name}: tarife satırı yok.")
    return tarife, uy


def indirim_oku(yol: Path) -> tuple[dict, list[dict]]:
    tablo, uy = {}, []
    for r in kayitlar(yol, INDIRIM_SUTUN, ("kod", "oran")):
        kod = metin(r["kod"]).upper()
        if not kod:
            continue
        oran, tur = para(r["oran"]), katla(r["tur"])
        if oran is None or oran < 0 or not (tur.startswith("indirim") or tur.startswith("surprim")):
            uy.append(uyari("Yüksek", "İndirim / sürprim satırı okunamadı", kod, f"{yol.name} satır {r['_satir']}: tür 'İndirim' veya 'Sürprim', oran pozitif sayı olmalı. Satır kullanılmadı."))
            continue
        t = [katla(x) for x in re.split(r"[|,;]", metin(r["teminat"])) if katla(x)]
        tablo[kod] = {"ad": metin(r["ad"]) or kod, "indirim": tur.startswith("indirim"), "oran": oran, "teminat": None if not t or "*" in metin(r["teminat"]) else set(t)}
    return tablo, uy


def riskleri_oku(yol: Path) -> tuple[list[dict], list[dict]]:
    riskler, uy = [], []
    for r in kayitlar(yol, RISK_SUTUN, ("no", "teminat", "bedel")):
        no, tem = metin(r["no"]), metin(r["teminat"])
        if not no or not tem:
            continue
        kim, bedel = f"{no} · {tem}", para(r["bedel"])
        if bedel is None or bedel <= 0:
            uy.append(uyari("Yüksek", "Bedel okunamadı", kim, f"{yol.name} satır {r['_satir']}: '{metin(r['bedel'])}'. Prim hesaplanmadı."))
            continue
        bas, bit, gun = tarih(r["bas"]), tarih(r["bit"]), para(r["gun"])
        if bas and bit:
            if bit <= bas:
                uy.append(uyari("Yüksek", "Bitiş tarihi başlangıçtan önce", kim, f"{bas:%d.%m.%Y} – {bit:%d.%m.%Y}. Prim hesaplanmadı."))
                continue
            tam_yil = (bit.day, bit.month, bit.year) == (bas.day, bas.month, bas.year + 1)
            gun = Decimal(365) if tam_yil else Decimal((bit - bas).days)
        elif (metin(r["bas"]) and not bas) or (metin(r["bit"]) and not bit):
            uy.append(uyari("Orta", "Tarih okunamadı", kim, "Başlangıç veya bitiş tarihi okunamadı; süre 365 gün sayıldı."))
            gun = None
        if gun is None or gun <= 0:
            gun = Decimal(365)
        riskler.append({"no": no, "sigortali": metin(r["sigortali"]), "sinif": metin(r["sinif"]), "teminat": tem, "bedel": bedel, "gun": gun,
                        "kodlar": [k.strip().upper() for k in re.split(r"[|,;]", metin(r["kodlar"])) if k.strip()]})
    return riskler, uy


def prim_hesapla(risk: dict, tarife: dict, indirimler: dict, zincirleme: bool = True, azami_indirim: Decimal | None = None) -> dict:
    """Tek teminat satırı için prim. Hesaplanamazsa 'hata' alanı dolar."""
    s = dict(risk) | {"fiyat": None, "baz": None, "katsayi": None, "carpan": None, "hesaplanan": None, "net": None, "asgari": False, "tavan": False, "uygulanan": [], "notlar": [], "hata": ""}
    t = tarife.get((katla(risk["teminat"]), katla(risk["sinif"]))) or tarife.get((katla(risk["teminat"]), "*"))
    if not t:
        s["hata"] = f"Tarifede '{risk['teminat']}' teminatı için '{risk['sinif'] or 'boş'}' risk sınıfı yok."
        return s
    carpan, toplam = BIR, Z
    for kod in dict.fromkeys(risk["kodlar"]):
        i = indirimler.get(kod)
        if not i:
            s["notlar"].append(f"'{kod}' kodu indirim / sürprim tablosunda yok; uygulanmadı.")
            continue
        if i["teminat"] is not None and katla(risk["teminat"]) not in i["teminat"]:
            s["notlar"].append(f"{kod} ({i['ad']}) bu teminata uygulanmaz; atlandı.")
            continue
        o = i["oran"] / 100 * (-1 if i["indirim"] else 1)
        carpan, toplam = carpan * (1 + o), toplam + o
        s["uygulanan"].append(f"{kod} {'−' if i['indirim'] else '+'}%{i['oran'].normalize():f}")
    carpan = carpan if zincirleme else 1 + toplam
    if azami_indirim is not None and carpan < 1 - azami_indirim / 100:
        s["notlar"].append(f"Toplam indirim %{yuvarla((1 - carpan) * 100)}; azami indirim tavanı %{azami_indirim.normalize():f} uygulandı.")
        carpan, s["tavan"] = 1 - azami_indirim / 100, True
    carpan = max(carpan, Z)
    s["fiyat"], s["katsayi"], s["carpan"] = t["fiyat"], risk["gun"] / 365, carpan
    s["baz"] = yuvarla(risk["bedel"] * t["fiyat"] / 1000)
    s["hesaplanan"] = yuvarla(risk["bedel"] * t["fiyat"] / 1000 * s["katsayi"] * carpan)
    s["net"] = s["hesaplanan"]
    if t["asgari"] is not None and s["net"] < t["asgari"]:
        s["net"], s["asgari"] = yuvarla(t["asgari"]), True
        s["notlar"].append(f"Hesaplanan prim {s['hesaplanan']:f}, asgari primin ({t['asgari']:f}) altında; asgari prim uygulandı.")
    return s


def hesapla(riskler: list[dict], tarife: dict, indirimler: dict, zincirleme: bool = True, azami_indirim: Decimal | None = None, vergi: Decimal = Z) -> dict:
    satirlar = [prim_hesapla(r, tarife, indirimler, zincirleme, azami_indirim) for r in riskler]
    teklifler, uy = {}, []
    for s in satirlar:
        t = teklifler.setdefault(s["no"], {"no": s["no"], "sigortali": s["sigortali"], "teminat": 0, "hesaplanamayan": 0, "net": Z})
        t["teminat"] += 1
        t["sigortali"] = t["sigortali"] or s["sigortali"]
        if s["hata"]:
            t["hesaplanamayan"] += 1
            uy.append(uyari("Yüksek", "Prim hesaplanamadı", f"{s['no']} · {s['teminat']}", s["hata"]))
        else:
            t["net"] += s["net"]
        for n in s["notlar"]:
            uy.append(uyari("Orta" if "tablosunda yok" in n else "Bilgi", "İndirim / sürprim notu", f"{s['no']} · {s['teminat']}", n))
    for t in teklifler.values():
        t["vergi"] = yuvarla(t["net"] * vergi / 100)
        t["brut"] = t["net"] + t["vergi"]
    return {"satirlar": satirlar, "teklifler": list(teklifler.values()), "uyarilar": uy, "vergi": vergi, "zincirleme": zincirleme}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Teklif Özeti"
    _baslik(ws, ["Teklif No", "Sigortalı", "Teminat Sayısı", "Hesaplanamayan", "Net Prim", f"Vergi (%{r['vergi'].normalize():f})", "Brüt Prim"], (14, 34, 10, 12, 16, 14, 16))
    for t in r["teklifler"]:
        _satir(ws, [t["no"], t["sigortali"], t["teminat"], t["hesaplanamayan"] or None, t["net"], t["vergi"], t["brut"]], para_s=(5, 6, 7), dolgu={4: KIRMIZI if t["hesaplanamayan"] else None})
    _satir(ws, ["Toplam", "", sum(t["teminat"] for t in r["teklifler"]), sum(t["hesaplanamayan"] for t in r["teklifler"]) or None, sum((t["net"] for t in r["teklifler"]), Z),
                sum((t["vergi"] for t in r["teklifler"]), Z), sum((t["brut"] for t in r["teklifler"]), Z)], para_s=(5, 6, 7))
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.append([])
    ws.append([f"İndirim / sürprim yöntemi: {'zincirleme (çarpımsal)' if r['zincirleme'] else 'toplamsal'}. Hesaplanamayan teminatlar toplama girmez."])
    ws = wb.create_sheet("Teminat Primleri")
    _baslik(ws, ["Teklif No", "Sigortalı", "Risk Sınıfı", "Teminat", "Bedel", "Fiyat (‰)", "Baz Prim (yıllık)", "Gün", "Süre Katsayısı", "İndirim / Sürprim", "Çarpan",
                 "Hesaplanan Prim", "Net Prim", "Not"], (14, 30, 10, 22, 18, 9, 16, 7, 10, 28, 9, 16, 16, 70))
    for s in r["satirlar"]:
        _satir(ws, [s["no"], s["sigortali"], s["sinif"], s["teminat"], s["bedel"], s["fiyat"], s["baz"], int(s["gun"]), yuvarla(s["katsayi"], "0.0001") if s["katsayi"] is not None else None,
                    ", ".join(s["uygulanan"]), yuvarla(s["carpan"], "0.0001") if s["carpan"] is not None else None, s["hesaplanan"], s["net"], s["hata"] or " ".join(s["notlar"])],
               para_s=(5, 7, 12, 13), dolgu={13: KIRMIZI if s["hata"] else SARI if s["asgari"] or s["tavan"] else None}, sar=(14,))
        for j in (9, 11):
            ws.cell(ws.max_row, j).number_format = "0.0000"
    uyari_sayfasi(wb, sirala(r["uyarilar"]))
    kaydet(wb, cikti)


def calistir(risk_yol: Path, tarife_yol: Path, indirim_yol: Path | None, cikti: Path, **secenek) -> dict:
    tarife, uy1 = tarife_oku(tarife_yol)
    indirimler, uy2 = indirim_oku(indirim_yol) if indirim_yol else ({}, [])
    riskler, uy3 = riskleri_oku(risk_yol)
    if not riskler:
        raise ValueError(f"{risk_yol.name}: hesaplanacak teminat satırı yok.")
    r = hesapla(riskler, tarife, indirimler, **secenek)
    r["uyarilar"] = sirala(uy1 + uy2 + uy3 + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Tarife fiyatı, bedel, risk sınıfı ve indirim-sürprim tablosundan teminat bazında prim hesaplar.")
    p.add_argument("--riskler", type=Path, default=ORNEK / "riskler.csv", help="Teklif No, Sigortalı, Risk Sınıfı, Teminat, Bedel, Başlangıç, Bitiş, İndirim / Sürprim")
    p.add_argument("--tarife", type=Path, default=ORNEK / "tarife.csv", help="Teminat, Risk Sınıfı (* = hepsi), Fiyat (‰), Asgari Prim")
    p.add_argument("--indirimler", type=Path, help="Kod, Ad, Tür (İndirim / Sürprim), Oran %%, Teminat (* = hepsi)")
    p.add_argument("--yontem", choices=("zincirleme", "toplamsal"), default="zincirleme", help="İndirim-sürprimlerin birleştirilme yöntemi (varsayılan zincirleme)")
    p.add_argument("--azami-indirim", help="Toplam indirim tavanı %% (varsayılan: tavan yok)")
    p.add_argument("--vergi", default="0", help="Net prime eklenecek vergi oranı %% (varsayılan 0)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "prim_hesabi.xlsx")
    a = p.parse_args(argv)
    indirim = a.indirimler or (ORNEK / "indirim_surprim.csv" if a.tarife == ORNEK / "tarife.csv" else None)
    for y in (a.riskler, a.tarife, indirim):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    vergi, tavan = para(a.vergi), para(a.azami_indirim) if a.azami_indirim else None
    if vergi is None or vergi < 0 or (a.azami_indirim and (tavan is None or not 0 <= tavan <= 100)):
        print("[X] --vergi sıfır veya pozitif; --azami-indirim 0-100 arasında olmalı")
        return 2
    try:
        r = calistir(a.riskler, a.tarife, indirim, a.cikti, zincirleme=a.yontem == "zincirleme", azami_indirim=tavan, vergi=vergi)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    for t in r["teklifler"]:
        print(f"[OK] {t['no']}: net prim {t['net']:,.2f}" + (f" · brüt {t['brut']:,.2f}" if vergi else "") + (f" · {t['hesaplanamayan']} teminat hesaplanamadı" if t["hesaplanamayan"] else ""))
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['kim']} · {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
