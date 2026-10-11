import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
TARIFE = {("yangin", "a"): {"fiyat": D("0.60"), "asgari": D(500)}, ("yangin", "b"): {"fiyat": D("0.90"), "asgari": D(500)}, ("deprem", "*"): {"fiyat": D("1.50"), "asgari": None},
          ("hirsizlik", "*"): {"fiyat": D("1.20"), "asgari": D(300)}}
IND = {"SPR": {"ad": "Sprinkler", "indirim": True, "oran": D(15), "teminat": {"yangin"}}, "HSZ": {"ad": "Hasarsızlık", "indirim": True, "oran": D(10), "teminat": None},
       "HSR": {"ad": "Hasar", "indirim": False, "oran": D(20), "teminat": None}}


def risk(teminat="Yangın", bedel=65_000_000, sinif="B", kodlar=(), gun=365, no="T1"):
    return {"no": no, "sigortali": "X", "sinif": sinif, "teminat": teminat, "bedel": D(bedel), "gun": D(gun), "kodlar": list(kodlar)}


class PrimTesti(unittest.TestCase):
    def test_baz_prim(self):
        s = main.prim_hesapla(risk(), TARIFE, IND)
        self.assertEqual((s["baz"], s["net"], s["carpan"], s["katsayi"], s["hata"]), (D("58500.00"), D("58500.00"), D(1), D(1), ""))     # 65 milyon × 0,90 ‰
        self.assertEqual(main.prim_hesapla(risk("Deprem", sinif="Z"), TARIFE, IND)["net"], D("97500.00"))                               # * = tüm sınıflar
        s = main.prim_hesapla(risk(sinif="D"), TARIFE, IND)
        self.assertEqual((s["net"], "risk sınıfı yok" in s["hata"]), (None, True))

    def test_indirim_surprim(self):
        s = main.prim_hesapla(risk(kodlar=["SPR", "HSZ"]), TARIFE, IND)
        self.assertEqual((s["net"], s["carpan"], s["uygulanan"]), (D("44752.50"), D("0.7650"), ["SPR −%15", "HSZ −%10"]))               # 58.500 × 0,85 × 0,90
        self.assertEqual(main.prim_hesapla(risk(kodlar=["SPR", "HSZ"]), TARIFE, IND, zincirleme=False)["net"], D("43875.00"))             # 58.500 × 0,75
        self.assertEqual(main.prim_hesapla(risk(kodlar=["HSR", "HSZ", "HSZ"]), TARIFE, IND)["net"], D("63180.00"))                        # × 1,20 × 0,90; tekrar eden kod bir kez
        s = main.prim_hesapla(risk("Deprem", kodlar=["SPR", "XYZ"]), TARIFE, IND)
        self.assertEqual((s["net"], s["uygulanan"], len(s["notlar"])), (D("97500.00"), [], 2))
        s = main.prim_hesapla(risk(kodlar=["SPR", "HSZ"]), TARIFE, IND, azami_indirim=D(20))
        self.assertEqual((s["net"], s["tavan"]), (D("46800.00"), True))
        self.assertEqual(main.prim_hesapla(risk(kodlar=["SPR"]), TARIFE, IND, azami_indirim=D(15))["tavan"], False)                       # sınır dahil

    def test_sure_ve_asgari(self):
        s = main.prim_hesapla(risk(bedel=10_000_000, sinif="A", gun=181), TARIFE, IND)
        self.assertEqual((s["baz"], s["net"], s["asgari"]), (D("6000.00"), D("2975.34"), False))                                         # 6.000 × 181 ÷ 365
        s = main.prim_hesapla(risk(bedel=500_000, sinif="A", gun=181), TARIFE, IND)
        self.assertEqual((s["hesaplanan"], s["net"], s["asgari"]), (D("148.77"), D("500.00"), True))
        self.assertEqual(main.prim_hesapla(risk("Deprem", bedel=1000), TARIFE, IND)["net"], D("1.50"))                                    # asgari primi olmayan teminat

    def test_teklif_toplami_ve_vergi(self):
        r = main.hesapla([risk(kodlar=["SPR", "HSZ"]), risk("Deprem", kodlar=["HSZ"]), risk(sinif="D", no="T2"), risk("Deprem", 5_000_000, no="T2")], TARIFE, IND, vergi=D(5))
        t1, t2 = r["teklifler"]
        self.assertEqual((t1["net"], t1["vergi"], t1["brut"], t1["hesaplanamayan"]), (D("132502.50"), D("6625.13"), D("139127.63"), 0))
        self.assertEqual((t2["net"], t2["hesaplanamayan"], [u["tur"] for u in r["uyarilar"]]), (D("7500.00"), 1, ["Prim hesaplanamadı"]))


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "p.xlsx"
            r = main.calistir(ORNEK / "riskler.csv", ORNEK / "tarife.csv", ORNEK / "indirim_surprim.csv", cikti)
            self.assertEqual({t["no"]: (t["net"], t["hesaplanamayan"]) for t in r["teklifler"]},
                             {"TK-4001": (D("142222.50"), 0), "TK-4002": (D("72000.00"), 0), "TK-4003": (D("800.00"), 0), "TK-4004": (D("7500.00"), 1), "TK-4005": (D("13770.00"), 0)})
            self.assertEqual([s["net"] for s in r["satirlar"] if s["no"] == "TK-4001"], [D("44752.50"), D("87750.00"), D("9720.00")])
            self.assertEqual(sorted((u["onem"], u["kim"]) for u in r["uyarilar"]),
                             [("Bilgi", "TK-4003 · Hırsızlık"), ("Bilgi", "TK-4003 · Yangın"), ("Bilgi", "TK-4005 · Yangın"), ("Orta", "TK-4004 · Deprem"), ("Yüksek", "TK-4004 · Yangın")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Teklif Özeti", "Teminat Primleri", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            r, t, i = (Path(tmp) / x for x in ("r.csv", "t.csv", "i.csv"))
            t.write_text("Teminat;Risk Sınıfı;Fiyat (‰)\nYangın;A;0,60\nYangın;a;0,70\nDeprem;;pahalı\nCam;;2\n", encoding="utf-8")
            i.write_text("Kod;Tür;Oran %;Teminat\nspr;İndirim;15;Yangın, Cam\nX;Bonus;5;\nY;Sürprim;;\n", encoding="utf-8")
            r.write_text("Teklif No;Risk Sınıfı;Teminat;Bedel;Başlangıç;Bitiş;İndirim / Sürprim\nT1;A;Yangın;1.000.000;29.02.2028;28.02.2029;spr\nT1;A;Cam;10.000;01.03.2028;01.03.2029;\n"
                         "T1;A;Deprem;çok;;;\nT2;A;Yangın;100;05.01.2027;01.01.2027;\nT3;A;Yangın;100;bozuk;;\n", encoding="utf-8")
            tar, uy = main.tarife_oku(t)
            self.assertEqual((sorted(tar), sorted(u["tur"] for u in uy)), ([("cam", "*"), ("yangin", "a")], ["Tarife fiyatı okunamadı", "Tarife satırı tekrar ediyor"]))
            ind, uy = main.indirim_oku(i)
            self.assertEqual((ind["SPR"]["teminat"], list(ind), len(uy)), ({"yangin", "cam"}, ["SPR"], 2))
            rs, uy = main.riskleri_oku(r)
            self.assertEqual([(x["no"], x["teminat"], x["gun"], x["kodlar"]) for x in rs], [("T1", "Yangın", D(365), ["SPR"]), ("T1", "Cam", D(365), []), ("T3", "Yangın", D(365), [])])
            self.assertEqual(sorted(u["tur"] for u in uy), ["Bedel okunamadı", "Bitiş tarihi başlangıçtan önce", "Tarih okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--riskler", str(r), "--tarife", str(t), "--indirimler", str(i), "--cikti", c]
            self.assertEqual(main.main(arg + ["--yontem", "toplamsal", "--vergi", "5", "--azami-indirim", "30"]), 0)
            self.assertEqual(main.main(arg + ["--vergi", "-1"]), 2)
            self.assertEqual(main.main(arg + ["--azami-indirim", "120"]), 2)
            self.assertEqual(main.main(["--riskler", str(Path(tmp) / "yok.csv")]), 1)
            t.write_text("Teminat;Fiyat\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
