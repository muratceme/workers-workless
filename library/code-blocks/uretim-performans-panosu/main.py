"""
Üretim Performans Panosu — Workers / Workless kod bloğu
Üretim › Üretim Müdürü

Günlük üretim kayıtlarından OEE, plan uyumu, fire ve iş kazası göstergelerini hat ve hafta bazında tek panoda toplar.
  Çalışma süresi       = planlı süre − duruş
  Kullanılabilirlik %  = Σ çalışma süresi ÷ Σ planlı süre × 100
  Performans %         = Σ (ideal çevrim süresi × üretilen adet) ÷ Σ çalışma süresi × 100
  Kalite %             = Σ sağlam adet ÷ Σ üretilen adet × 100
  OEE %                = kullanılabilirlik × performans × kalite
  Plan uyumu %         = Σ sağlam adet ÷ Σ planlanan adet × 100
  Fire %               = Σ (üretilen − sağlam) ÷ Σ üretilen × 100
  Hafta                = tarihin ISO hafta numarası (pazartesi başlar)
Oranlar günlük oranların ortalaması değil, toplamların oranıdır. Hedefler (--oee 60, --plan 90, --fire 3) sizin politikanızdır.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"tarih": ("tarih", "uretim tarihi"), "hat": ("hat", "makine", "uretim hatti", "is merkezi"), "planli": ("planli sure dk", "planli sure", "planli uretim suresi"), "durus": ("durus dk", "durus", "durus suresi"),
         "cevrim": ("ideal cevrim suresi sn", "ideal cevrim suresi", "cevrim suresi sn", "cevrim suresi"), "uretilen": ("uretilen adet", "uretilen", "toplam uretim"), "saglam": ("saglam adet", "saglam", "kabul adet"),
         "fire": ("fire adet", "fire", "hurda adet", "ret adet"), "plan": ("planlanan adet", "plan adet", "hedef adet"), "kaza": ("is kazasi", "is kazasi sayisi", "kaza")}
TOPLA = ("planli", "calisma", "ideal", "uretilen", "saglam", "plan", "kaza")


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("tarih", "hat", "planli", "cevrim", "uretilen")):
        hat = metin(r["hat"])
        if not hat:
            continue
        kim = f"{hat} · {metin(r['tarih'])}"
        t, d = tarih(r["tarih"]), {a: para(r[a]) for a in ("planli", "durus", "cevrim", "uretilen", "saglam", "fire", "plan", "kaza")}
        durus = d["durus"] or Decimal(0)
        saglam = d["saglam"] if d["saglam"] is not None else d["uretilen"] - (d["fire"] or 0) if d["uretilen"] is not None else None
        hata = (t is None or d["planli"] is None or d["planli"] <= 0 or d["cevrim"] is None or d["cevrim"] <= 0 or d["uretilen"] is None or d["uretilen"] < 0 or not 0 <= durus <= d["planli"] or not 0 <= saglam <= d["uretilen"]
                or any(metin(r[a]) and d[a] is None for a in d))
        if hata:
            uy.append(uyari("Yüksek", "Satır okunamadı", kim, f"Satır {r['_satir']}: tarih ya da sayılardan biri okunamadı veya tutarsız (duruş planlı süreden, sağlam adet üretilenden büyük olamaz). Satır alınmadı."))
            continue
        if d["saglam"] is None and d["fire"] is None:
            uy.append(uyari("Bilgi", "Kalite verisi yok", kim, f"Satır {r['_satir']}: sağlam adet de fire adet de yazılmamış; fire sıfır sayıldı."))
        l.append({"tarih": t, "hat": hat, "planli": d["planli"], "durus": durus, "cevrim": d["cevrim"], "uretilen": d["uretilen"], "saglam": saglam, "plan": d["plan"] or Decimal(0), "kaza": int(d["kaza"] or 0)})
    return l, uy


def _yuzde(pay, payda) -> Decimal | None:
    return yuvarla(Decimal(pay) * 100 / Decimal(payda)) if payda else None


def gosterge(l: list[dict]) -> dict:
    t = {a: sum((Decimal(s[a]) for s in l), Decimal(0)) for a in TOPLA}
    oee = yuvarla(t["ideal"] / t["planli"] * t["saglam"] / t["uretilen"] * 100) if t["planli"] and t["uretilen"] and t["calisma"] else None
    return {**t, "gun": len(l), "kullan": _yuzde(t["calisma"], t["planli"]), "perf": _yuzde(t["ideal"], t["calisma"]), "kalite": _yuzde(t["saglam"], t["uretilen"]), "oee": oee, "uyum": _yuzde(t["saglam"], t["plan"]),
            "fire": t["uretilen"] - t["saglam"], "fire_oran": _yuzde(t["uretilen"] - t["saglam"], t["uretilen"]), "kaza": int(t["kaza"])}


def pano(kayit: list[dict], oee: Decimal = Decimal(60), plan: Decimal = Decimal(90), fire: Decimal = Decimal(3)) -> dict:
    if not kayit:
        raise ValueError("Panoya girecek üretim kaydı yok.")
    for s in kayit:
        s["calisma"] = s["planli"] - s["durus"]
        s["ideal"] = s["cevrim"] * s["uretilen"] / 60
        iso = s["tarih"].isocalendar()
        s["hafta"] = f"{iso[0]}-H{iso[1]:02d}"

    def kirilim(*alanlar: str) -> list[dict]:
        grup = {}
        for s in kayit:
            grup.setdefault(tuple(katla(s[a]) for a in alanlar), []).append(s)
        return [{**{a: g[0][a] for a in alanlar}, "ilk": min(s["tarih"] for s in g), "son": max(s["tarih"] for s in g), **gosterge(g)} for _, g in sorted(grup.items())]

    detay, uy = kirilim("hat", "hafta"), []
    for d in detay:
        kim = f"{d['hat']} · {d['hafta']}"
        if d["kaza"]:
            uy.append(uyari("Yüksek", "İş kazası", kim, f"{d['ilk']:%d.%m.%Y}–{d['son']:%d.%m.%Y} arasında {d['kaza']} iş kazası kaydı var."))
        if d["oee"] is not None and d["oee"] < oee:
            uy.append(uyari("Orta", "OEE hedefin altında", kim, f"OEE %{d['oee']} (hedef %{oee}): kullanılabilirlik %{d['kullan']}, performans %{d['perf']}, kalite %{d['kalite']}."))
        if d["perf"] is not None and d["perf"] > 100:
            uy.append(uyari("Orta", "Performans %100'ün üstünde", kim, f"Performans %{d['perf']}. İdeal çevrim süresi gerçek kapasiteden uzun tanımlanmış olabilir; değeri kontrol edin."))
        if d["uyum"] is not None and d["uyum"] < plan:
            uy.append(uyari("Orta", "Plan uyumu düşük", kim, f"Sağlam üretim {d['saglam']:,.0f}, plan {d['plan']:,.0f} (%{d['uyum']}; hedef %{plan})."))
        if d["fire_oran"] is not None and d["fire_oran"] > fire:
            uy.append(uyari("Orta", "Fire oranı yüksek", kim, f"Fire {d['fire']:,.0f} adet, üretilen {d['uretilen']:,.0f} (%{d['fire_oran']}; sınır %{fire})."))
    return {"kayit": sorted(kayit, key=lambda s: (s["tarih"], katla(s["hat"]))), "hatlar": kirilim("hat"), "haftalar": kirilim("hafta"), "detay": detay, "toplam": gosterge(kayit), "uyarilar": uy}


BASLIK = ["Gün", "OEE %", "Kullanılabilirlik %", "Performans %", "Kalite %", "Plan Uyumu %", "Fire %", "Üretilen", "Sağlam", "Fire", "Planlanan", "Planlı Süre (dk)", "Çalışma Süresi (dk)", "İş Kazası"]


def _deger(g: dict) -> list:
    return [g["gun"], g["oee"], g["kullan"], g["perf"], g["kalite"], g["uyum"], g["fire_oran"], g["uretilen"], g["saglam"], g["fire"], g["plan"], g["planli"], g["calisma"], g["kaza"] or None]


def rapor_yaz(cikti: Path, r: dict, oee: Decimal, plan: Decimal, fire: Decimal) -> None:
    wb = Workbook()

    def renk(g: dict, n: int) -> dict:
        return {n + 2: None if g["oee"] is None else KIRMIZI if g["oee"] < oee else YESIL, n + 6: KIRMIZI if g["uyum"] is not None and g["uyum"] < plan else None, n + 7: KIRMIZI if g["fire_oran"] is not None and g["fire_oran"] > fire else None,
                n + 14: KIRMIZI if g["kaza"] else None}

    for i, (baslik, alan, adlar) in enumerate((("Pano (Hat × Hafta)", "detay", {"hat": "Hat", "hafta": "Hafta"}), ("Hatlar", "hatlar", {"hat": "Hat"}), ("Haftalar", "haftalar", {"hafta": "Hafta"}))):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title, n = baslik, len(adlar)
        _baslik(ws, list(adlar.values()) + BASLIK, (18,) * n + (6, 9, 13, 12, 9, 11, 8, 11, 11, 9, 11, 11, 12, 8), get_column_letter(n + 1) + "2")
        for g in r[alan]:
            _satir(ws, [g[a] for a in adlar] + _deger(g), para_s=range(n + 2, n + 8), dolgu=renk(g, n))
        if n == 1:
            k = _satir(ws, ["TOPLAM"] + _deger(r["toplam"]), para_s=range(3, 9))
            for h in ws[k]:
                h.font = Font(bold=True)
        ws.append([])
        ws.append([f"Hedefler: OEE en az %{oee}, plan uyumu en az %{plan}, fire en çok %{fire}. Oranlar toplamlardan hesaplanır."])
    ws = wb.create_sheet("Günlük Kayıtlar")
    _baslik(ws, ["Tarih", "Hafta", "Hat", "Planlı Süre (dk)", "Duruş (dk)", "İdeal Çevrim (sn)", "Üretilen", "Sağlam", "Planlanan", "İş Kazası", "OEE %"], (12, 10, 18, 11, 9, 11, 10, 10, 10, 8, 9))
    for s in r["kayit"]:
        _satir(ws, [s["tarih"], s["hafta"], s["hat"], s["planli"], s["durus"], s["cevrim"], s["uretilen"], s["saglam"], s["plan"] or None, s["kaza"] or None, gosterge([s])["oee"]], para_s=(11,), tarih_s=(1,))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, oee: Decimal = Decimal(60), plan: Decimal = Decimal(90), fire: Decimal = Decimal(3)) -> dict:
    kayit, uy = oku(yol)
    r = pano(kayit, oee, plan, fire)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, oee, plan, fire)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="OEE, plan uyumu, fire ve iş kazası göstergelerini hat ve hafta bazında tek panoda toplar.")
    p.add_argument("--kayitlar", type=Path, help="Tarih, Hat, Planlı Süre (dk), Duruş (dk), İdeal Çevrim Süresi (sn), Üretilen Adet, Sağlam Adet ya da Fire Adet, Planlanan Adet, İş Kazası")
    p.add_argument("--oee", default="60", help="OEE hedefi, yüzde; altı uyarılır (varsayılan 60)")
    p.add_argument("--plan", default="90", help="Plan uyumu hedefi, yüzde; altı uyarılır (varsayılan 90)")
    p.add_argument("--fire", default="3", help="Fire oranı sınırı, yüzde; üstü uyarılır (varsayılan 3)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "uretim_performans_panosu.xlsx")
    a = p.parse_args(argv)
    yol = a.kayitlar or ORNEK / "gunluk_uretim.csv"
    oee, plan, fire = para(a.oee), para(a.plan), para(a.fire)
    if any(x is None or not 0 <= x <= 100 for x in (oee, plan, fire)):
        print("[X] --oee, --plan ve --fire 0 ile 100 arasında yüzde olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, oee, plan, fire)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {len(r['hatlar'])} hat · {len(r['haftalar'])} hafta · OEE %{t['oee']} (kullanılabilirlik %{t['kullan']}, performans %{t['perf']}, kalite %{t['kalite']})"
          + (f" · plan uyumu %{t['uyum']}" if t["uyum"] is not None else "") + f" · fire %{t['fire_oran']} · iş kazası {t['kaza']}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
