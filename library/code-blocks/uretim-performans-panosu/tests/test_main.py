import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def s(hat, t, planli, durus, cevrim, uretilen, saglam, plan=0, kaza=0):
    return {"tarih": t, "hat": hat, "planli": D(planli), "durus": D(durus), "cevrim": D(cevrim), "uretilen": D(uretilen), "saglam": D(saglam), "plan": D(plan), "kaza": kaza}


def veri():
    return [s("A", date(2026, 10, 5), 480, 60, 30, 700, 665, 800), s("a", date(2026, 10, 12), 480, 30, 30, 810, 800, 800, 1), s("B", date(2026, 10, 5), 480, 0, 60, 500, 500, 450)]


def k(g):
    return (g["gun"], g["oee"], g["kullan"], g["perf"], g["kalite"], g["uyum"], g["fire_oran"], g["fire"], g["kaza"])


class PanoTesti(unittest.TestCase):
    def test_gostergeler(self):
        r = main.pano(veri())
        self.assertEqual([(g["hat"], g["hafta"]) + k(g) for g in r["detay"]],
                         [("A", "2026-H41", 1, D("69.27"), D("87.50"), D("83.33"), D("95.00"), D("83.13"), D("5.00"), D(35), 0), ("a", "2026-H42", 1, D("83.33"), D("93.75"), D("90.00"), D("98.77"), D("100.00"), D("1.23"), D(10), 1),
                          ("B", "2026-H41", 1, D("104.17"), D("100.00"), D("104.17"), D("100.00"), D("111.11"), D("0.00"), D(0), 0)])
        self.assertEqual([(g["hat"],) + k(g) for g in r["hatlar"]],
                         [("A", 2, D("76.30"), D("90.63"), D("86.78"), D("97.02"), D("91.56"), D("2.98"), D(45), 1), ("B", 1, D("104.17"), D("100.00"), D("104.17"), D("100.00"), D("111.11"), D("0.00"), D(0), 0)])
        self.assertEqual([(g["hafta"],) + k(g) for g in r["haftalar"]],
                         [("2026-H41", 2, D("85.96"), D("93.75"), D("94.44"), D("97.08"), D("93.20"), D("2.92"), D(35), 0), ("2026-H42", 1, D("83.33"), D("93.75"), D("90.00"), D("98.77"), D("100.00"), D("1.23"), D(10), 1)])
        self.assertEqual(k(r["toplam"]), (3, D("85.20"), D("93.75"), D("92.96"), D("97.76"), D("95.85"), D("2.24"), D(45), 1))
        self.assertEqual((r["toplam"]["planli"], r["toplam"]["calisma"], r["toplam"]["ideal"], r["toplam"]["uretilen"], r["toplam"]["saglam"], r["toplam"]["plan"]), (D(1440), D(1350), D(1255), D(2010), D(1965), D(2050)))
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Orta", "Fire oranı yüksek", "A · 2026-H41"), ("Orta", "Performans %100'ün üstünde", "B · 2026-H41"), ("Orta", "Plan uyumu düşük", "A · 2026-H41"), ("Yüksek", "İş kazası", "a · 2026-H42")])

    def test_hedefler_ve_hafta(self):
        u = lambda **kw: sorted(x["tur"] for x in main.pano([s("A", date(2026, 10, 5), 480, 60, 30, 700, 665, 800)], **kw)["uyarilar"])   # noqa: E731
        self.assertEqual((u(oee=D("69.27"), plan=D("83.13"), fire=D(5)), u(oee=D("69.28"), plan=D("83.14"), fire=D("4.99"))), ([], ["Fire oranı yüksek", "OEE hedefin altında", "Plan uyumu düşük"]))
        l = [s("A", t, 100, 0, 60, 100, 100) for t in (date(2025, 12, 28), date(2025, 12, 29), date(2026, 1, 4), date(2026, 1, 5), date(2027, 1, 3))]
        r = main.pano(l)
        self.assertEqual(([(g["hafta"], g["gun"], g["ilk"], g["son"]) for g in r["haftalar"]], r["toplam"]["uyum"], r["toplam"]["oee"]),
                         ([("2025-H52", 1, date(2025, 12, 28), date(2025, 12, 28)), ("2026-H01", 2, date(2025, 12, 29), date(2026, 1, 4)), ("2026-H02", 1, date(2026, 1, 5), date(2026, 1, 5)), ("2026-H53", 1, date(2027, 1, 3), date(2027, 1, 3))],
                          None, D("100.00")))
        g = main.pano([s("A", date(2026, 10, 5), 100, 100, 60, 0, 0)])["toplam"]
        self.assertEqual((g["oee"], g["kullan"], g["perf"], g["kalite"], g["fire_oran"]), (None, D("0.00"), None, None, None))
        with self.assertRaises(ValueError):
            main.pano([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "gunluk_uretim.csv", cikti)
            self.assertEqual([(g["hat"], g["hafta"]) + k(g) for g in r["detay"]],
                             [("Hat 1", "2026-H41", 3, D("78.00"), D("90.00"), D("88.89"), D("97.50"), D("92.37"), D("2.50"), D(54), 0), ("Hat 1", "2026-H42", 3, D("63.33"), D("80.00"), D("83.33"), D("95.00"), D("75.00"), D("5.00"), D(90), 1),
                              ("Hat 2", "2026-H41", 3, D("85.67"), D("93.33"), D("92.86"), D("98.85"), D("95.19"), D("1.15"), D(9), 0), ("Hat 2", "2026-H42", 3, D("55.33"), D("66.67"), D("85.00"), D("97.65"), D("61.48"), D("2.35"), D(12), 0)])
            self.assertEqual([(g["hat"],) + k(g) for g in r["hatlar"]],
                             [("Hat 1", 6, D("70.67"), D("85.00"), D("86.27"), D("96.36"), D("83.68"), D("3.64"), D(144), 1), ("Hat 2", 6, D("70.50"), D("80.00"), D("89.58"), D("98.37"), D("78.33"), D("1.63"), D(21), 0)])
            self.assertEqual(k(r["toplam"]), (12, D("70.22"), D("82.50"), D("87.88"), D("96.86"), D("82.28"), D("3.14"), D(165), 1))
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Yüksek", "İş kazası", "Hat 1 · 2026-H42"), ("Orta", "Fire oranı yüksek", "Hat 1 · 2026-H42"), ("Orta", "OEE hedefin altında", "Hat 2 · 2026-H42"), ("Orta", "Plan uyumu düşük", "Hat 1 · 2026-H42"),
                              ("Orta", "Plan uyumu düşük", "Hat 2 · 2026-H42")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Pano (Hat × Hafta)", "Hatlar", "Haftalar", "Günlük Kayıtlar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "u.csv"
            y.write_text("Üretim Tarihi;Makine;Planlı Süre;Duruş;Çevrim Süresi;Üretilen;Fire;Hedef Adet;Kaza\n05.10.2026;M1;480;30,5;12,5;1.000;40;1.100;1\n06.10.2026;M1;480;;10;500;;;\nx;M1;480;;10;500;;;\n07.10.2026;M1;480;500;10;500;;;\n"
                         "07.10.2026;M1;480;;10;500;501;;\n07.10.2026;M1;0;;10;500;;;\n07.10.2026;M1;480;;;500;;;\n07.10.2026;;480;;10;500;;;\n07.10.2026;M1;480;;10;500;;y;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["tarih"], x["hat"], x["planli"], x["durus"], x["cevrim"], x["uretilen"], x["saglam"], x["plan"], x["kaza"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([(date(2026, 10, 5), "M1", D(480), D("30.5"), D("12.5"), D(1000), D(960), D(1100), 1), (date(2026, 10, 6), "M1", D(480), D(0), D(10), D(500), D(500), D(0), 0)],
                              [("Kalite verisi yok", "M1 · 06.10.2026"), ("Satır okunamadı", "M1 · x")] + [("Satır okunamadı", "M1 · 07.10.2026")] * 5))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--kayitlar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--oee", "70", "--plan", "95", "--fire", "2,5"]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--oee", "101"], ["--plan", "-1"], ["--fire", "x"])], [2] * 3)
            self.assertEqual(main.main(["--kayitlar", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Tarih;Hat;Planlı Süre (dk);Üretilen Adet\n05.10.2026;A;480;5\n", "Tarih;Hat;Planlı Süre (dk);İdeal Çevrim Süresi (sn);Üretilen Adet\n05.10.2026;;480;10;5\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
