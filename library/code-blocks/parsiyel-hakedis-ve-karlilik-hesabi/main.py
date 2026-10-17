"""
Parsiyel Hakediş ve Kârlılık Hesabı — Gerisi Bende kod bloğu
Parsiyel Operasyon › Parsiyel Operasyon Uzmanı

Araçtaki yüklerin navlun gelirlerini, araç maliyetini ve yük bazında kâr payını hesaplar.
  Yükün maliyet payı = araç maliyeti × yükün ölçüsü ÷ araçtaki yüklerin toplam ölçüsü     (ölçü: --anahtar ldm | kg | gelir; varsayılan ldm)
                       kuruş farkı aracın son yüküne eklenir; payların toplamı araç maliyetine eşittir
  Yük kârı           = navlun geliri − maliyet payı;   marj % = kâr ÷ navlun geliri × 100
  Araç kârı          = araçtaki yüklerin navlun geliri toplamı − araç maliyeti
  Doluluk %          = toplam LDM ÷ araç kapasitesi (dosyadaki "Kapasite LDM", yoksa --kapasite-ldm, varsayılan 13,6)
  Dolu LDM maliyeti  = araç maliyeti ÷ toplam LDM
Araç maliyeti, alt nakliyeciye ödenecek hakediş veya öz mal aracın sefer maliyetidir; siz verirsiniz. Boş kalan kapasitenin maliyeti araçtaki yüklere dağılır.
Dağıtım anahtarı bir tercihtir; anahtar değişince yük bazındaki kâr değişir, araç kârı değişmez. Tek para birimi varsayılır.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


YUK = {"sefer": ("sefer", "sefer no", "arac", "plaka", "arac sefer no"), "musteri": ("musteri", "gonderici", "cari"), "ref": ("yuk", "yuk ref", "referans", "pozisyon", "dosya no"),
       "gelir": ("navlun geliri", "navlun", "gelir", "satis"), "ldm": ("ldm", "yukleme metresi"), "kg": ("agirlik", "kg", "agirlik kg", "brut agirlik")}
SEFER = {"sefer": YUK["sefer"], "maliyet": ("arac maliyeti", "maliyet", "hakedis", "alt nakliyeci bedeli", "sefer maliyeti"), "kapasite": ("kapasite ldm", "kapasite", "yukleme uzunlugu")}
ANAHTAR = {"ldm": "LDM", "kg": "ağırlık", "gelir": "navlun geliri"}
Z = Decimal(0)


def oku(yuk_yol: Path, sefer_yol: Path) -> tuple[list[dict], dict, list[dict]]:
    uy, yukler, seferler = [], [], {}
    for r in kayitlar(yuk_yol, YUK, ("sefer", "gelir")):
        s = metin(r["sefer"])
        if not s:
            continue
        kim = f"{s} · {metin(r['ref']) or metin(r['musteri']) or 'yük'} (satır {r['_satir']})"
        gelir = para(r["gelir"])
        if gelir is None or gelir < 0:
            uy.append(uyari("Yüksek", "Navlun geliri okunamadı", kim, f"'{metin(r['gelir'])}'. Yük alınmadı; araçtaki diğer yüklerin maliyet payı olduğundan yüksek çıkar."))
            continue
        x = {"sefer": s, "musteri": metin(r["musteri"]) or "Belirtilmemiş", "ref": metin(r["ref"]), "gelir": gelir}
        for a in ("ldm", "kg"):
            v = para(r[a])
            if v is None and metin(r[a]) or (v is not None and v < 0):
                uy.append(uyari("Orta", "Ölçü okunamadı", kim, f"{a.upper()} '{metin(r[a])}'. 0 sayıldı."))
                v = None
            x[a] = v or Z
        yukler.append(x)
    for r in kayitlar(sefer_yol, SEFER, ("sefer", "maliyet")):
        s, m, k = metin(r["sefer"]), para(r["maliyet"]), para(r["kapasite"])
        if not s:
            continue
        if m is None or m < 0:
            uy.append(uyari("Yüksek", "Araç maliyeti okunamadı", s, f"{sefer_yol.name} satır {r['_satir']}: '{metin(r['maliyet'])}'. Bu aracın kârlılığı hesaplanmadı."))
            continue
        if katla(s) in seferler:
            uy.append(uyari("Yüksek", "Mükerrer sefer", s, f"{sefer_yol.name}: aynı sefer birden fazla satırda; yalnız ilki alındı."))
            continue
        seferler[katla(s)] = {"maliyet": m, "kapasite": k if k and k > 0 else None}
    return yukler, seferler, uy


def hesapla(yukler: list[dict], seferler: dict, anahtar: str = "ldm", kapasite: Decimal = Decimal("13.6")) -> dict:
    if not yukler:
        raise ValueError("Hesaplanacak yük yok.")
    grup, uy = defaultdict(list), []
    for y in yukler:
        grup[katla(y["sefer"])].append(y)
    ozet, hesaplanan = [], []
    for k, liste in grup.items():
        s = seferler.get(k)
        if s is None:
            uy.append(uyari("Yüksek", "Aracın maliyeti yok", liste[0]["sefer"], f"{len(liste)} yük var ama sefer maliyeti dosyasında bu sefer yok. Kârlılık hesaplanmadı."))
            continue
        kullanilan, toplam = anahtar, sum((y[anahtar] for y in liste), Z)
        if toplam == 0:
            kullanilan, toplam = "gelir", sum((y["gelir"] for y in liste), Z)
            uy.append(uyari("Orta", "Dağıtım anahtarı değişti", liste[0]["sefer"], f"Yüklerin {ANAHTAR[anahtar]} toplamı 0; maliyet navlun gelirine göre dağıtıldı."))
        elif any(y[anahtar] == 0 for y in liste):
            uy.append(uyari("Orta", "Ölçüsü olmayan yük", liste[0]["sefer"], f"{sum(y[anahtar] == 0 for y in liste)} yükün {ANAHTAR[anahtar]} değeri 0; bu yüklere maliyet payı düşmedi."))
        kalan = s["maliyet"]
        for i, y in enumerate(liste):
            y["oran"] = yuvarla(y[kullanilan] / toplam * 100) if toplam else Z
            y["pay"] = kalan if i == len(liste) - 1 else (yuvarla(s["maliyet"] * y[kullanilan] / toplam) if toplam else Z)
            kalan -= y["pay"]
            y["kar"] = y["gelir"] - y["pay"]
            y["marj"] = yuvarla(y["kar"] / y["gelir"] * 100) if y["gelir"] else None
            hesaplanan.append(y)
        gelir, ldm, kap = sum((y["gelir"] for y in liste), Z), sum((y["ldm"] for y in liste), Z), s["kapasite"] or kapasite
        ozet.append({"sefer": liste[0]["sefer"], "yuk": len(liste), "gelir": gelir, "maliyet": s["maliyet"], "kar": gelir - s["maliyet"], "marj": yuvarla((gelir - s["maliyet"]) / gelir * 100) if gelir else None,
                     "ldm": ldm, "kg": sum((y["kg"] for y in liste), Z), "kapasite": kap, "doluluk": yuvarla(ldm / kap * 100) if ldm else None, "bos": max(kap - ldm, Z) if ldm else None,
                     "ldm_maliyet": yuvarla(s["maliyet"] / ldm) if ldm else None, "anahtar": ANAHTAR[kullanilan], "zarar_yuk": sum(y["kar"] < 0 for y in liste)})
    if not ozet:
        raise ValueError("Hiçbir aracın maliyeti bulunamadı: yük dosyasındaki sefer numaraları sefer maliyeti dosyasıyla eşleşmiyor.")
    mus = {}
    for y in hesaplanan:
        m = mus.setdefault(katla(y["musteri"]), {"musteri": y["musteri"], "yuk": 0, "gelir": Z, "pay": Z})
        m["yuk"] += 1
        m["gelir"] += y["gelir"]
        m["pay"] += y["pay"]
    for m in mus.values():
        m["kar"] = m["gelir"] - m["pay"]
        m["marj"] = yuvarla(m["kar"] / m["gelir"] * 100) if m["gelir"] else None
    g, c = sum((o["gelir"] for o in ozet), Z), sum((o["maliyet"] for o in ozet), Z)
    return {"yukler": sorted(hesaplanan, key=lambda y: (y["sefer"], y["kar"])), "seferler": sorted(ozet, key=lambda o: (o["kar"], o["sefer"])), "musteri": sorted(mus.values(), key=lambda m: m["kar"]),
            "genel": {"gelir": g, "maliyet": c, "kar": g - c, "marj": yuvarla((g - c) / g * 100) if g else None, "yuk": len(hesaplanan), "zarar_yuk": sum(y["kar"] < 0 for y in hesaplanan),
                      "zarar_sefer": sum(o["kar"] < 0 for o in ozet)}, "uyarilar": uy, "maliyetsiz": sorted(k for k in seferler if k not in grup)}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Araç Özeti"
    _baslik(ws, ["Sefer / Araç", "Yük Sayısı", "Navlun Geliri", "Araç Maliyeti (hakediş)", "Kâr", "Marj %", "Toplam LDM", "Kapasite LDM", "Doluluk %", "Boş LDM", "Dolu LDM Başına Maliyet", "Toplam Ağırlık (kg)",
                 "Dağıtım Anahtarı", "Zarar Eden Yük"], (18, 8, 14, 14, 14, 9, 10, 10, 10, 9, 13, 12, 14, 9))
    for o in r["seferler"]:
        _satir(ws, [o["sefer"], o["yuk"], o["gelir"], o["maliyet"], o["kar"], o["marj"], o["ldm"], o["kapasite"], o["doluluk"], o["bos"], o["ldm_maliyet"], o["kg"], o["anahtar"], o["zarar_yuk"] or None],
               para_s=(3, 4, 5, 6, 7, 8, 9, 10, 11), dolgu={5: KIRMIZI if o["kar"] < 0 else YESIL})
    g = r["genel"]
    _satir(ws, ["Toplam", g["yuk"], g["gelir"], g["maliyet"], g["kar"], g["marj"]], para_s=(3, 4, 5, 6))
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.append([])
    ws.append(["Yük bazındaki kâr dağıtım anahtarına bağlıdır; araç kârı anahtardan bağımsızdır. Boş kapasitenin maliyeti araçtaki yüklere dağılır."])
    ws = wb.create_sheet("Yük Bazında")
    _baslik(ws, ["Sefer / Araç", "Müşteri", "Yük Ref", "Navlun Geliri", "LDM", "Ağırlık (kg)", "Pay Oranı %", "Maliyet Payı", "Kâr", "Marj %"], (18, 26, 16, 14, 9, 12, 10, 14, 14, 9))
    for y in r["yukler"]:
        _satir(ws, [y["sefer"], y["musteri"], y["ref"], y["gelir"], y["ldm"], y["kg"], y["oran"], y["pay"], y["kar"], y["marj"]], para_s=(4, 5, 7, 8, 9, 10), dolgu={9: KIRMIZI if y["kar"] < 0 else None})
    ws = wb.create_sheet("Müşteri Bazında")
    _baslik(ws, ["Müşteri", "Yük Sayısı", "Navlun Geliri", "Maliyet Payı", "Kâr", "Marj %"], (28, 8, 14, 14, 14, 9))
    for m in r["musteri"]:
        _satir(ws, [m["musteri"], m["yuk"], m["gelir"], m["pay"], m["kar"], m["marj"]], para_s=(3, 4, 5, 6), dolgu={5: KIRMIZI if m["kar"] < 0 else None})
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yuk_yol: Path, sefer_yol: Path, cikti: Path, anahtar: str = "ldm", kapasite: Decimal = Decimal("13.6")) -> dict:
    yukler, seferler, uy = oku(yuk_yol, sefer_yol)
    r = hesapla(yukler, seferler, anahtar, kapasite)
    uy += r["uyarilar"]
    for o in r["seferler"]:
        if o["doluluk"] is not None and o["doluluk"] > 100:
            uy.append(uyari("Yüksek", "Doluluk %100'ün üzerinde", o["sefer"], f"Toplam {o['ldm']:f} LDM, kapasite {o['kapasite']:f} LDM. LDM değerlerini veya kapasiteyi kontrol edin."))
    for k in r["maliyetsiz"]:
        uy.append(uyari("Orta", "Maliyeti olan aracın yükü yok", k, "Sefer maliyeti dosyasında var, yük dosyasında bu seferle yük yok."))
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Parsiyel araçlarda yük bazında maliyet payı ve kârlılık hesaplar.")
    p.add_argument("--yukler", type=Path, default=ORNEK / "yukler.csv", help="Sefer, Müşteri, Yük Ref, Navlun Geliri, LDM, Ağırlık (kg)")
    p.add_argument("--seferler", type=Path, default=ORNEK / "sefer_maliyetleri.csv", help="Sefer, Araç Maliyeti, Kapasite LDM")
    p.add_argument("--anahtar", choices=tuple(ANAHTAR), default="ldm", help="Araç maliyetinin yüklere dağıtım ölçüsü (varsayılan ldm)")
    p.add_argument("--kapasite-ldm", default="13,6", help="Dosyada kapasitesi yazmayan araçlar için kapasite (varsayılan 13,6 LDM)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "parsiyel_karlilik.xlsx")
    a = p.parse_args(argv)
    k = para(a.kapasite_ldm)
    if k is None or k <= 0:
        print("[X] --kapasite-ldm pozitif sayı olmalı")
        return 2
    for y in (a.yukler, a.seferler):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.yukler, a.seferler, a.cikti, a.anahtar, k)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    g = r["genel"]
    print(f"[OK] {len(r['seferler'])} araç · {g['yuk']} yük · gelir {g['gelir']:,.2f} · maliyet {g['maliyet']:,.2f} · kâr {g['kar']:,.2f} · {g['zarar_sefer']} zarar eden araç, {g['zarar_yuk']} zarar eden yük")
    for o in r["seferler"]:
        print(f"[{'!' if o['kar'] < 0 else 'OK'}] {o['sefer']}: kâr {o['kar']:,.2f}" + (f" · doluluk %{o['doluluk']}" if o["doluluk"] is not None else ""))
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
