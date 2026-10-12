import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def yk(gelir, ldm, kg=0, sefer="S", musteri="M"):
    return {"sefer": sefer, "musteri": musteri, "ref": "", "gelir": D(gelir), "ldm": D(ldm), "kg": D(kg)}


def sm(maliyet, kapasite=None):
    return {"s": {"maliyet": D(maliyet), "kapasite": None if kapasite is None else D(kapasite)}}


class HesapTesti(unittest.TestCase):
    def test_dagitim(self):
        l = [yk(1500, 4, 3000), yk(1200, "3.2", 5000), yk(900, "4.8", 2000)]
        r = main.hesapla(l, sm(3000))
        self.assertEqual([(y["oran"], y["pay"], y["kar"], y["marj"]) for y in sorted(r["yukler"], key=lambda y: -y["gelir"])],
                         [(D("33.33"), D("1000.00"), D("500.00"), D("33.33")), (D("26.67"), D("800.00"), D("400.00"), D("33.33")), (D("40.00"), D("1200.00"), D("-300.00"), D("-33.33"))])
        o = r["seferler"][0]
        self.assertEqual((o["gelir"], o["kar"], o["marj"], o["ldm"], o["doluluk"], o["bos"], o["ldm_maliyet"], o["zarar_yuk"]), (D(3600), D(600), D("16.67"), D("12.0"), D("88.24"), D("1.6"), D("250.00"), 1))
        r = main.hesapla([yk(1500, 4, 3000), yk(1200, "3.2", 5000), yk(900, "4.8", 2000)], sm(3000), "kg")
        self.assertEqual([y["pay"] for y in sorted(r["yukler"], key=lambda y: -y["gelir"])], [D("900.00"), D("1500.00"), D("600.00")])
        r = main.hesapla([yk(1500, 4), yk(1200, 3), yk(900, 5)], sm(3000), "gelir")
        self.assertEqual({y["marj"] for y in r["yukler"]}, {D("16.67")})                                       # gelire göre dağıtımda marjlar eşit
        r = main.hesapla([yk(100, 1), yk(100, 1), yk(100, 1)], sm(100))
        self.assertEqual((sorted(y["pay"] for y in r["yukler"]), sum(y["pay"] for y in r["yukler"])), ([D("33.33"), D("33.33"), D("33.34")], D(100)))   # kuruş farkı son yükte

    def test_eksik_veri_ve_ozet(self):
        r = main.hesapla([yk(500, 0), yk(700, 0)], sm(600))
        self.assertEqual(([y["pay"] for y in r["yukler"]], r["seferler"][0]["anahtar"], r["seferler"][0]["doluluk"], [u["tur"] for u in r["uyarilar"]]),
                         ([D("250.00"), D("350.00")], "navlun geliri", None, ["Dağıtım anahtarı değişti"]))
        r = main.hesapla([yk(500, 0), yk(700, 2)], sm(600))
        self.assertEqual((sorted(y["pay"] for y in r["yukler"]), [u["tur"] for u in r["uyarilar"]]), ([D(0), D("600.00")], ["Ölçüsü olmayan yük"]))
        l = [yk(1000, 3, sefer="P1", musteri="A"), yk(1100, 3, sefer="P1", musteri="B"), yk(2000, 10, sefer="P2", musteri="a"), yk(5, 1, sefer="P3")]
        r = main.hesapla(l, {"p1": {"maliyet": D(2800), "kapasite": None}, "p2": {"maliyet": D(1500), "kapasite": D(8)}, "p9": {"maliyet": D(1), "kapasite": None}})
        self.assertEqual([(o["sefer"], o["kar"], o["kapasite"], o["doluluk"]) for o in r["seferler"]], [("P1", D(-700), D("13.6"), D("44.12")), ("P2", D(500), D(8), D("125.00"))])
        self.assertEqual([(m["musteri"], m["yuk"], m["gelir"], m["kar"]) for m in r["musteri"]], [("B", 1, D(1100), D(-300)), ("A", 2, D(3000), D(100))])
        self.assertEqual((r["genel"], [u["tur"] for u in r["uyarilar"]], r["maliyetsiz"]),
                         ({"gelir": D(4100), "maliyet": D(4300), "kar": D(-200), "marj": D("-4.88"), "yuk": 3, "zarar_yuk": 2, "zarar_sefer": 1}, ["Aracın maliyeti yok"], ["p9"]))
        for y, s in (([], sm(1)), ([yk(1, 1)], {})):
            with self.assertRaises(ValueError):
                main.hesapla(y, s)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "p.xlsx"
            r = main.calistir(ORNEK / "yukler.csv", ORNEK / "sefer_maliyetleri.csv", cikti)
            self.assertEqual([(o["sefer"], o["gelir"], o["kar"], o["doluluk"]) for o in r["seferler"]], [("P-02", D(2100), D(-700), D("44.12")), ("P-01", D(3600), D(600), D("88.24"))])
            self.assertEqual({y["ref"]: y["kar"] for y in r["yukler"]}, {"Y-101": D(500), "Y-102": D(400), "Y-103": D(-300), "Y-201": D(-400), "Y-202": D(-300)})
            self.assertEqual((r["genel"]["kar"], r["genel"]["marj"], r["uyarilar"], load_workbook(cikti).sheetnames), (D(-100), D("-1.75"), [], ["Araç Özeti", "Yük Bazında", "Müşteri Bazında", "Uyarılar"]))

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, s = Path(tmp) / "y.csv", Path(tmp) / "s.csv"
            y.write_text("Sefer;Müşteri;Navlun;LDM;Ağırlık\nP1;A;1.250,50;4,5;ağır\nP1;B;yok;1;1\n;C;1;1;1\n", encoding="utf-8")
            s.write_text("Sefer;Hakediş;Kapasite\nP1;2.000;7,2\np1;1;\nP2;pahalı;\n", encoding="utf-8")
            yl, sl, uy = main.oku(y, s)
            self.assertEqual(([(x["sefer"], x["gelir"], x["ldm"], x["kg"]) for x in yl], sl), ([("P1", D("1250.50"), D("4.5"), D(0))], {"p1": {"maliyet": D(2000), "kapasite": D("7.2")}}))
            self.assertEqual(sorted(u["tur"] for u in uy), ["Araç maliyeti okunamadı", "Mükerrer sefer", "Navlun geliri okunamadı", "Ölçü okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--yukler", str(y), "--seferler", str(s), "--cikti", c]
            self.assertEqual(main.main(arg + ["--anahtar", "kg", "--kapasite-ldm", "7,2"]), 0)
            self.assertEqual(main.main(arg + ["--kapasite-ldm", "0"]), 2)
            self.assertEqual(main.main(arg[:3] + [str(Path(tmp) / "yok.csv")]), 1)
            s.write_text("Sefer;Araç Maliyeti\nZ9;5\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
