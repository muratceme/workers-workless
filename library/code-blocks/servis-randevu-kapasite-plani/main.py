"""
Servis Randevu Kapasite Planı — Gerisi Bende kod bloğu
Satış Sonrası Hizmetler (Servis) › Servis Danışmanı
Teknisyen kapasitesi ve iş türü sürelerine göre günlük randevu kapasitesini, doluluğu ve eklenebilecek iş sayısını hesaplar.
  Mevcut saat     = o gün çalışan (izinli olmayan) teknisyenlerin günlük saatleri toplamı
  Kapasite (saat) = Σ günlük saat × verimlilik % × (1 − --rezerv %)   (rezerv: randevusuz gelen, garanti, tekrar işler için ayrılan pay)
  Planlanan       = randevuların süresi (satırda Süre yoksa iş türünün standart süresi; "+" ile birden fazla iş)
  Doluluk %       = planlanan ÷ kapasite × 100;  eklenebilir iş = ⌊kalan saat ÷ iş türü süresi⌋
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
TEKNISYEN = {"ad": ("teknisyen", "ad soyad", "personel", "usta"), "saat": ("gunluk saat", "calisma saati", "saat"), "verim": ("verimlilik", "verimlilik orani", "verim"),
             "gunler": ("calisma gunleri", "gunler", "calistigi gunler"), "izin": ("izin", "izinli gunler", "izin tarihleri")}
ISTURU = {"ad": ("is turu", "is", "islem", "hizmet"), "sure": ("standart sure", "standart sure saat", "sure saat", "sure", "saat", "standart saat")}
RANDEVU = {"tarih": ("tarih", "randevu tarihi", "gun"), "plaka": ("plaka", "arac"), "is": ("is turu", "is", "islem", "hizmet"), "sure": ("sure", "planlanan sure", "saat"),
           "musteri": ("musteri", "ad soyad")}
GUN_AD = ("Pzt", "Sal", "Çar", "Per", "Cum", "Cmt", "Paz")
SIFIR = Decimal(0)


def _gun(p: str) -> int | None:
    if p.startswith("pazartesi") or p in ("pzt", "pts", "pt"):
        return 0
    if p.startswith("cumartesi") or p in ("cmt", "cts", "ct"):
        return 5
    for i, onek in ((6, "paz"), (1, "sal"), (2, "car"), (3, "per"), (4, "cum")):
        if p.startswith(onek):
            return i
    return None


def gun_listesi(x) -> set[int] | None:
    """'Pzt-Cmt', 'Pazartesi, Çarşamba', 'Paz' → hafta günü numaraları (0 = pazartesi); boş → tüm günler."""
    s = katla(str(x or "").replace("-", " aralik ").replace("–", " aralik "))
    if not s:
        return set(range(7))
    sonuc, parca = set(), s.split()
    i = 0
    while i < len(parca):
        a = _gun(parca[i])
        if a is None:
            return None
        if i + 2 < len(parca) and parca[i + 1] == "aralik":
            b = _gun(parca[i + 2])
            if b is None:
                return None
            sonuc |= set(range(a, b + 1)) if a <= b else set(range(a, 7)) | set(range(0, b + 1))
            i += 3
        else:
            sonuc.add(a)
            i += 1
    return sonuc


def tarih_listesi(x) -> set[date] | None:
    s = metin(x)
    if not s:
        return set()
    sonuc = set()
    for p in re.split(r"[,;]+", s):
        p = p.strip()
        if not p:
            continue
        m = re.fullmatch(r"(\S+)\s*[-–]\s*(\S+)", p) if p.count(".") > 2 or "/" in p else None
        if m:
            a, b = tarih(m.group(1)), tarih(m.group(2))
            if a is None or b is None or b < a:
                return None
            sonuc |= {a + timedelta(days=i) for i in range((b - a).days + 1)}
        else:
            t = tarih(p)
            if t is None:
                return None
            sonuc.add(t)
    return sonuc


def teknisyen_oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, TEKNISYEN, ("ad", "saat")):
        ad = metin(r["ad"])
        if not ad:
            continue
        saat, verim = para(r["saat"]), para(r["verim"]) if metin(r["verim"]) else Decimal(100)
        gunler, izin = gun_listesi(r["gunler"]), tarih_listesi(r["izin"])
        if saat is None or not 0 < saat <= 24 or verim is None or not 0 < verim <= 200 or gunler is None or izin is None:
            uy.append(uyari("Yüksek", "Teknisyen satırı okunamadı", ad, f"Satır {r['_satir']}: günlük saat, verimlilik, çalışma günleri ('{metin(r['gunler'])}') ya da izin tarihleri okunamadı."))
            continue
        l.append({"ad": ad, "saat": saat, "verim": verim, "gunler": gunler, "izin": izin})
    return l, uy


def is_oku(yol: Path) -> tuple[dict, list[dict]]:
    uy, d = [], {}
    for r in kayitlar(yol, ISTURU, ("ad", "sure")):
        ad = metin(r["ad"])
        if not ad:
            continue
        s = para(r["sure"])
        if s is None or s <= 0:
            uy.append(uyari("Yüksek", "İş türü okunamadı", ad, f"Satır {r['_satir']}: standart süre okunamadı."))
            continue
        d[katla(ad)] = {"ad": ad, "sure": s}
    return d, uy


def randevu_oku(yol: Path, isler: dict) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, RANDEVU, ("tarih", "is")):
        if not metin(r["tarih"]) and not metin(r["is"]):
            continue
        g = tarih(r["tarih"])
        kim = metin(r["plaka"]) or f"Satır {r['_satir']}"
        sure = para(r["sure"]) if metin(r["sure"]) else None
        if g is None or (metin(r["sure"]) and (sure is None or sure <= 0)):
            uy.append(uyari("Yüksek", "Randevu okunamadı", kim, f"Satır {r['_satir']}: tarih ya da süre okunamadı."))
            continue
        turler = [p.strip() for p in re.split(r"\s*\+\s*", metin(r["is"])) if p.strip()]
        bilinmeyen = [t for t in turler if katla(t) not in isler]
        if sure is None:
            if bilinmeyen:
                uy.append(uyari("Yüksek", "İş türü tanımsız", kim, f"{g:%d.%m.%Y}: '{', '.join(bilinmeyen)}' iş türü listesinde yok ve süre yazılmamış; randevu kapasiteye girmedi."))
                continue
            sure = sum((isler[katla(t)]["sure"] for t in turler), SIFIR)
        l.append({"tarih": g, "plaka": metin(r["plaka"]), "is": " + ".join(turler), "sure": sure, "musteri": metin(r["musteri"])})
    return l, uy


def planla(teknisyenler: list[dict], isler: dict, randevular: list[dict], baslangic: date, gun: int, kapali: set[int] | None = None, rezerv: Decimal = SIFIR,
           hedef: Decimal = Decimal(85)) -> dict:
    if not teknisyenler:
        raise ValueError("Okunabilen teknisyen yok.")
    kapali = kapali or set()
    uy, gunler = [], []
    son = baslangic + timedelta(days=gun - 1)
    for i in range(gun):
        g = baslangic + timedelta(days=i)
        r_gun = [r for r in randevular if r["tarih"] == g]
        planlanan = sum((r["sure"] for r in r_gun), SIFIR)
        acik = g.weekday() not in kapali
        calisan = [t for t in teknisyenler if acik and g.weekday() in t["gunler"] and g not in t["izin"]]
        mevcut = sum((t["saat"] for t in calisan), SIFIR)
        kapasite = yuvarla(sum((t["saat"] * t["verim"] / 100 for t in calisan), SIFIR) * (1 - rezerv / 100), "0.1")
        doluluk = yuvarla(planlanan * 100 / kapasite) if kapasite else None
        kalan = kapasite - planlanan
        etiket = f"{g:%d.%m.%Y} {GUN_AD[g.weekday()]}"
        if not acik:
            if r_gun:
                uy.append(uyari("Orta", "Kapalı güne randevu", etiket, f"{len(r_gun)} randevu ({planlanan} saat) servisin kapalı olduğu güne yazılmış."))
            gunler.append({"tarih": g, "etiket": etiket, "acik": False, "teknisyen": 0, "izinli": 0, "mevcut": SIFIR, "kapasite": SIFIR, "randevu": len(r_gun), "planlanan": planlanan,
                           "doluluk": None, "kalan": None, "eklenebilir": {}})
            continue
        if not calisan and r_gun:
            uy.append(uyari("Yüksek", "Çalışan teknisyen yok", etiket, f"{len(r_gun)} randevu var ama o gün çalışan teknisyen yok."))
        elif doluluk is not None and doluluk > 100:
            uy.append(uyari("Yüksek", "Kapasite aşıldı", etiket, f"Planlanan {planlanan} saat, kapasite {kapasite} saat (%{doluluk}); {-kalan} saatlik iş başka güne kaydırılmalı."))
        elif doluluk is not None and doluluk > hedef:
            uy.append(uyari("Orta", "Hedef doluluk aşıldı", etiket, f"Doluluk %{doluluk} (hedef %{hedef}); kalan {kalan} saat."))
        elif doluluk is not None and doluluk < 50:
            uy.append(uyari("Bilgi", "Boş kapasite", etiket, f"Doluluk %{doluluk}; {kalan} saat boş. Bakım hatırlatması ya da kampanya aramasıyla doldurulabilir."))
        izinli = sum(1 for t in teknisyenler if g.weekday() in t["gunler"] and g in t["izin"])
        ekle = {i["ad"]: (int(kalan // i["sure"]) if kalan > 0 else 0) for i in isler.values()}
        gunler.append({"tarih": g, "etiket": etiket, "acik": True, "teknisyen": len(calisan), "izinli": izinli, "mevcut": mevcut, "kapasite": kapasite, "randevu": len(r_gun),
                       "planlanan": planlanan, "doluluk": doluluk, "kalan": kalan, "eklenebilir": ekle})
    disarida = [r for r in randevular if not baslangic <= r["tarih"] <= son]
    if disarida:
        uy.append(uyari("Bilgi", "Dönem dışı randevu", f"{len(disarida)} randevu", f"{baslangic:%d.%m.%Y}–{son:%d.%m.%Y} dışındaki randevular plana girmedi."))
    acik = [g for g in gunler if g["acik"]]
    kap = sum((g["kapasite"] for g in acik), SIFIR)
    pl = sum((g["planlanan"] for g in gunler), SIFIR)
    toplam = {"kapasite": kap, "planlanan": pl, "doluluk": yuvarla(pl * 100 / kap) if kap else None, "randevu": sum(g["randevu"] for g in gunler), "kalan": kap - pl}
    sira = sorted(randevular, key=lambda r: (r["tarih"], katla(r["plaka"])))
    return {"gunler": gunler, "toplam": toplam, "isler": sorted(isler.values(), key=lambda i: i["sure"]), "randevular": [r for r in sira if baslangic <= r["tarih"] <= son], "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, hedef: Decimal) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Günlük Kapasite"
    _baslik(ws, ["Gün", "Açık", "Çalışan Teknisyen", "İzinli", "Mevcut Saat", "Kapasite (saat)", "Randevu", "Planlanan (saat)", "Doluluk %", "Kalan (saat)"], (16, 6, 10, 7, 9, 10, 8, 10, 9, 10))
    for g in r["gunler"]:
        d = g["doluluk"]
        _satir(ws, [g["etiket"], "Evet" if g["acik"] else "Kapalı", g["teknisyen"], g["izinli"], g["mevcut"], g["kapasite"], g["randevu"], g["planlanan"], d, g["kalan"]], para_s=(5, 6, 8, 9, 10),
               dolgu={9: None if d is None else (KIRMIZI if d > 100 else SARI if d > hedef else MAVI if d < 50 else YESIL)})
    t = r["toplam"]
    n = _satir(ws, ["TOPLAM", "", "", "", "", t["kapasite"], t["randevu"], t["planlanan"], t["doluluk"], t["kalan"]], para_s=(6, 8, 9, 10))
    for h in ws[n]:
        h.font = Font(bold=True)
    ws = wb.create_sheet("Eklenebilir İş")
    _baslik(ws, ["Gün"] + [f"{i['ad']} ({i['sure']} sa)" for i in r["isler"]], (16,) + (12,) * len(r["isler"]))
    for g in r["gunler"]:
        if g["acik"]:
            _satir(ws, [g["etiket"]] + [g["eklenebilir"].get(i["ad"], 0) for i in r["isler"]])
    ws = wb.create_sheet("Randevular")
    _baslik(ws, ["Tarih", "Plaka", "Müşteri", "İş", "Süre (saat)"], (11, 12, 20, 40, 9))
    for x in r["randevular"]:
        _satir(ws, [x["tarih"], x["plaka"], x["musteri"], x["is"], x["sure"]], tarih_s=(1,), para_s=(5,))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(tek_yol: Path, is_yol: Path, rnd_yol: Path, cikti: Path, baslangic: date | None = None, gun: int = 6, kapali: set[int] | None = None, rezerv: Decimal = SIFIR,
             hedef: Decimal = Decimal(85)) -> dict:
    tek, uy1 = teknisyen_oku(tek_yol)
    isler, uy2 = is_oku(is_yol)
    rnd, uy3 = randevu_oku(rnd_yol, isler)
    if baslangic is None:
        if not rnd:
            raise ValueError("Randevu yok; --baslangic verin.")
        baslangic = min(r["tarih"] for r in rnd)
    r = planla(tek, isler, rnd, baslangic, gun, kapali, rezerv, hedef)
    r["uyarilar"] = sirala(uy1 + uy2 + uy3 + r["uyarilar"])
    rapor_yaz(cikti, r, hedef)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Teknisyen kapasitesi ve iş sürelerine göre günlük randevu kapasitesini ve boşlukları hesaplar.")
    p.add_argument("--teknisyenler", type=Path, help="Teknisyen, Günlük Saat, Verimlilik %%, Çalışma Günleri (ör. Pzt-Cmt), İzin (tarih ya da aralık listesi)")
    p.add_argument("--isler", type=Path, help="İş Türü, Standart Süre (saat)")
    p.add_argument("--randevular", type=Path, help="Tarih, Plaka, Müşteri, İş Türü (birden fazlaysa + ile), Süre (isteğe bağlı)")
    p.add_argument("--baslangic", help="Plan başlangıcı GG.AA.YYYY (varsayılan ilk randevu günü)")
    p.add_argument("--gun", type=int, default=6, help="Kaç gün planlanacak (varsayılan 6)")
    p.add_argument("--kapali", default="Paz", help="Servisin kapalı olduğu günler (varsayılan Paz; boş = her gün açık)")
    p.add_argument("--rezerv", default="0", help="Randevusuz / acil işler için kapasiteden ayrılan yüzde (varsayılan 0)")
    p.add_argument("--hedef", default="85", help="Hedef doluluk yüzdesi (varsayılan 85)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "servis_kapasite_plani.xlsx")
    a = p.parse_args(argv)
    yollar = [a.teknisyenler or ORNEK / "teknisyenler.csv", a.isler or ORNEK / "is_turleri.csv", a.randevular or ORNEK / "randevular.csv"]
    bas = tarih(a.baslangic) if a.baslangic else None
    kapali = gun_listesi(a.kapali) if a.kapali.strip() else set()
    rez, hed = para(a.rezerv), para(a.hedef)
    if (a.baslangic and bas is None) or kapali is None or len(kapali) == 7 or rez is None or not 0 <= rez < 100 or hed is None or not 0 < hed <= 150 or not 1 <= a.gun <= 62:
        print("[X] --baslangic GG.AA.YYYY, --kapali gün adları (Pzt…Paz), --rezerv 0-99, --hedef 1-150, --gun 1-62 olmalı")
        return 2
    for y in yollar:
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(*yollar, a.cikti, bas, a.gun, kapali, rez, hed)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {len(r['gunler'])} gün · kapasite {t['kapasite']} saat · planlanan {t['planlanan']} saat ({t['randevu']} randevu) · doluluk %{t['doluluk']}")
    for g in r["gunler"]:
        print(f"     {g['etiket']}: " + (f"%{g['doluluk']} · kalan {g['kalan']} saat" if g["acik"] and g["doluluk"] is not None else "kapalı / teknisyen yok"))
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Plan: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
