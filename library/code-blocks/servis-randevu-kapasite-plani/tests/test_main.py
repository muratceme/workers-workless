import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
PZT = date(2026, 10, 19)


def t(ad, saat, verim=100, gunler="Pzt-Cmt", izin=""):
    return {"ad": ad, "saat": D(saat), "verim": D(verim), "gunler": main.gun_listesi(gunler), "izin": main.tarih_listesi(izin)}


def r(g, sure, plaka="P"):
    return {"tarih": date(2026, 10, g), "plaka": plaka, "is": "x", "sure": D(sure), "musteri": ""}


ISLER = {"bakim": {"ad": "Bakım", "sure": D("1.5")}, "fren": {"ad": "Fren", "sure": D(1)}}


class PlanTesti(unittest.TestCase):
    def test_kapasite(self):
        tek = [t("A", "7.5", 120), t("B", "7.5", 80, izin="20.10.2026"), t("C", 9, 100, "Pzt-Cum")]
        rnd = [r(19, 10), r(19, 12), r(20, 20), r(21, 2), r(24, 1), r(25, 3), r(30, 1)]
        x = main.planla(tek, ISLER, rnd, PZT, 7, {6}, D(10), D(85))
        g = {a["tarih"].day: a for a in x["gunler"]}
        # Pzt: (9 + 6 + 9) × 0,9 = 21,6 saat; 22 saat planlı → %101,85
        self.assertEqual((g[19]["mevcut"], g[19]["kapasite"], g[19]["doluluk"], g[19]["kalan"], g[19]["eklenebilir"]), (D("24"), D("21.6"), D("101.85"), D("-0.4"), {"Bakım": 0, "Fren": 0}))
        # Sal: B izinli → (9 + 9) × 0,9 = 16,2; Cmt: C çalışmaz → (9 + 6) × 0,9 = 13,5
        self.assertEqual((g[20]["teknisyen"], g[20]["izinli"], g[20]["kapasite"], g[20]["doluluk"]), (2, 1, D("16.2"), D("123.46")))
        self.assertEqual((g[21]["doluluk"], g[21]["eklenebilir"]), (D("9.26"), {"Bakım": 13, "Fren": 19}))
        self.assertEqual((g[24]["kapasite"], g[25]["acik"], g[25]["kapasite"]), (D("13.5"), False, D(0)))
        tp = x["toplam"]
        self.assertEqual((tp["kapasite"], tp["planlanan"], tp["randevu"]), (D("116.1"), D(48), 6))
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"][:10]) for u in x["uyarilar"]),
                         [("Bilgi", "Boş kapasite", "21.10.2026"), ("Bilgi", "Boş kapasite", "22.10.2026"), ("Bilgi", "Boş kapasite", "23.10.2026"), ("Bilgi", "Boş kapasite", "24.10.2026"),
                          ("Bilgi", "Dönem dışı randevu", "1 randevu"), ("Orta", "Kapalı güne randevu", "25.10.2026"), ("Yüksek", "Kapasite aşıldı", "19.10.2026"),
                          ("Yüksek", "Kapasite aşıldı", "20.10.2026")])
        y = main.planla([t("A", 8, gunler="Cmt")], ISLER, [r(19, 1)], PZT, 1)
        self.assertEqual(y["uyarilar"][0]["tur"], "Çalışan teknisyen yok")
        with self.assertRaises(ValueError):
            main.planla([], ISLER, [], PZT, 1)

    def test_donusum(self):
        self.assertEqual([main.gun_listesi(x) for x in ("", "Pzt-Cum", "Cmt–Pzt", "Paz", "Pzt, Çar, Cuma", "x")],
                         [set(range(7)), {0, 1, 2, 3, 4}, {5, 6, 0}, {6}, {0, 2, 4}, None])
        self.assertEqual([main.gun_listesi(x) for x in ("Pazartesi-Cumartesi", "Cumartesi", "Pazar")], [{0, 1, 2, 3, 4, 5}, {5}, {6}])
        self.assertEqual(main.tarih_listesi("19.10.2026-21.10.2026, 25.10.2026"), {date(2026, 10, 19), date(2026, 10, 20), date(2026, 10, 21), date(2026, 10, 25)})
        self.assertEqual([main.tarih_listesi(x) for x in ("", "x", "21.10.2026-19.10.2026")], [set(), None, None])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            x = main.calistir(ORNEK / "teknisyenler.csv", ORNEK / "is_turleri.csv", ORNEK / "randevular.csv", cikti, kapali={6})
            tp = x["toplam"]
            self.assertEqual((tp["kapasite"], tp["planlanan"], tp["randevu"], tp["doluluk"]), (D("264.6"), D("193.0"), 112, D("72.94")))
            self.assertEqual([g["doluluk"] for g in x["gunler"]], [D("96.03"), D("71.73"), D("120.51"), D("73.48"), D("37.52"), D("36.69")])
            self.assertEqual([(u["onem"], u["tur"]) for u in x["uyarilar"]][:3], [("Yüksek", "Kapasite aşıldı"), ("Yüksek", "İş türü tanımsız"), ("Orta", "Hedef doluluk aşıldı")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Günlük Kapasite", "Eklenebilir İş", "Randevular", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            ty, iy, ry = Path(tmp) / "t.csv", Path(tmp) / "i.csv", Path(tmp) / "r.csv"
            ty.write_text("Usta;Çalışma Saati;Verim;Günler;İzin Tarihleri\nA;8;;;\nB;25;;;\nC;8;x;;\nD;8;;Pazartesi-Salı;\nE;8;;Hafta içi;\n", encoding="utf-8")
            iy.write_text("İşlem;Süre\nBakım;1,5\nFren;0\n", encoding="utf-8")
            ry.write_text("Randevu Tarihi;Plaka;İşlem;Süre\n19.10.2026;A;Bakım + bakım;\n19.10.2026;B;Kaporta;\n19.10.2026;C;Kaporta;3\nx;D;Bakım;\n20.10.2026;E;Bakım;0\n", encoding="utf-8")
            l, uy = main.teknisyen_oku(ty)
            self.assertEqual(([(x["ad"], x["verim"], len(x["gunler"])) for x in l], len(uy)), ([("A", D(100), 7), ("D", D(100), 2)], 3))
            isler, uy = main.is_oku(iy)
            self.assertEqual((list(isler), len(uy)), (["bakim"], 1))
            l, uy = main.randevu_oku(ry, isler)
            self.assertEqual(([(x["plaka"], x["sure"]) for x in l], [u["tur"] for u in uy]), ([("A", D(3)), ("C", D(3))], ["İş türü tanımsız", "Randevu okunamadı", "Randevu okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--teknisyenler", str(ty), "--isler", str(iy), "--randevular", str(ry), "--cikti", c]
            self.assertEqual(main.main(arg + ["--baslangic", "19.10.2026", "--gun", "7", "--kapali", "", "--rezerv", "10", "--hedef", "90"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--baslangic", "x"], ["--kapali", "Pzt-Paz"], ["--kapali", "x"], ["--rezerv", "100"], ["--hedef", "0"], ["--gun", "0"])], [2] * 6)
            self.assertEqual(main.main(["--teknisyenler", str(Path(tmp) / "yok.csv")]), 1)
            ry.write_text("Tarih;İşlem\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
