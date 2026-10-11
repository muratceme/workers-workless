"""
Günlük Likidite Raporu — Workers / Workless kod bloğu
Hazine ve Aktif-Pasif Yönetimi › Hazine Uzmanı

Vadesi gelen varlık ve yükümlülüklerin nakit akışlarından para birimi bazında günlük ve haftalık likidite açığı / fazlası
tablosunu hazırlar.
  Kullanılabilir hazır değer = tutar × (1 − iskonto %)                (nakit, serbest menkul kıymet vb.)
  Günlük net akış            = girişler − çıkışlar
  Gün sonu likidite          = kullanılabilir hazır değerler + kümülatif net akış
  Likidite açığı             : gün sonu likiditenin eksiye düştüğü gün ve tutar.
  Dayanma süresi             : gün sonu likiditenin ilk kez eksiye düştüğü güne kadar geçen takvim günü.
  Stres senaryosu            : "Beklenen" girişler --giris-kesinti % azaltılır, "Beklenen" çıkışlar --cikis-artis %
                               artırılır; "Kesin" akışlar değişmez. Hazır değerlere iskonto her iki senaryoda uygulanır.
  Hafta sonu ve --tatiller dosyasındaki günlere düşen akışlar bir sonraki iş gününe kaydırılır (uyarı verilir).
  Rapor tarihinden önceki akışlar "gecikmiş" sayılıp ilk güne alınır.
Para birimleri ayrı ayrı raporlanır; çevrim yapılmaz. Yasal likidite oranları (LKO vb.) bu paketin kapsamında değildir.
İnternete bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: TL ve USD, 30 gün
    python main.py --akislar akislar.xlsx --hazir hazir_degerler.xlsx --tarih 05.10.2026 --gun 30 --giris-kesinti 50 --cikis-artis 20
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


AKIS_SUTUNLARI = {"tarih": ("tarih", "vade", "vade tarihi", "valor"), "kalem": ("kalem", "aciklama", "islem"), "yon": ("yon", "tur", "giris cikis"),
                  "tutar": ("tutar",), "pb": ("para birimi", "pb"), "kesinlik": ("kesinlik", "durum"), "karsi": ("karsi taraf", "musteri", "banka")}
HAZIR_SUTUNLARI = {"kalem": ("kalem", "aciklama", "hazir deger"), "tutar": ("tutar", "bakiye"), "pb": ("para birimi", "pb"), "iskonto": ("iskonto", "iskonto yuzde", "kesinti")}


def bin_(x) -> str:
    return f"{x:,.0f}".replace(",", ".")


def pb_duzelt(x) -> str:
    return metin(x).upper().replace("TRY", "TL") or "TL"


def is_gunu_mu(d: date, tatiller: set[date]) -> bool:
    return d.weekday() < 5 and d not in tatiller


def sonraki_is_gunu(d: date, tatiller: set[date]) -> date:
    while not is_gunu_mu(d, tatiller):
        d += timedelta(days=1)
    return d


def oku(a_yol: Path, h_yol: Path | None, rapor: date, t_yol: Path | None = None) -> tuple[list[dict], list[dict], set[date], list[dict]]:
    uy, akislar, hazir, tatiller = [], [], [], set()
    if t_yol:
        tatiller = {t for t in (tarih(r[0]) for r in tablo_oku(t_yol)) if t}
    kaydirilan, gecikmis = 0, 0
    for r in kayitlar(a_yol, AKIS_SUTUNLARI, ("tarih", "yon", "tutar")):
        if not (metin(r["tarih"]) or metin(r["tutar"])):
            continue
        t, v, yon = tarih(r["tarih"]), para(r["tutar"]), katla(r["yon"])
        giris = yon.startswith(("giris", "tahsil", "alacak", "varlik")) or yon == "+"
        cikis = yon.startswith(("cikis", "odeme", "borc", "yukumluluk")) or yon == "-"
        if not t or v is None or v < 0 or not (giris or cikis):
            uy.append(uyari("Yüksek", "Akış satırı okunamadı", metin(r["kalem"]) or f"Satır {r['_satir']}",
                            f"Satır {r['_satir']}: tarih '{metin(r['tarih'])}', yön '{metin(r['yon'])}' (Giriş / Çıkış), tutar '{metin(r['tutar'])}'. Satır atlandı."))
            continue
        asil = t
        if t < rapor:
            gecikmis += 1
            t = rapor
        if not is_gunu_mu(t, tatiller):
            kaydirilan += 1
            t = sonraki_is_gunu(t, tatiller)
        akislar.append({"tarih": t, "asil": asil, "kalem": metin(r["kalem"]), "giris": giris, "tutar": v, "pb": pb_duzelt(r["pb"]),
                        "kesin": not katla(r["kesinlik"]).startswith(("beklenen", "tahmin", "olasi")), "karsi": metin(r["karsi"])})
    if gecikmis:
        uy.append(uyari("Orta", "Tarihi geçmiş akış", a_yol.name, f"{gecikmis} akışın tarihi rapor tarihinden önce; ilk güne alındı. Gerçekleşmiş akışları dosyadan çıkarın."))
    if kaydirilan:
        uy.append(uyari("Bilgi", "İş günü olmayan tarih", a_yol.name, f"{kaydirilan} akış hafta sonu veya tatile denk geliyor; bir sonraki iş gününe kaydırıldı."))
    if h_yol:
        for r in kayitlar(h_yol, HAZIR_SUTUNLARI, ("kalem", "tutar")):
            if not metin(r["kalem"]):
                continue
            v, i = para(r["tutar"]), para(r["iskonto"]) or Decimal(0)
            if v is None or v < 0 or not (0 <= i <= 100):
                uy.append(uyari("Yüksek", "Hazır değer okunamadı", r["kalem"], f"Satır {r['_satir']}: tutar '{metin(r['tutar'])}', iskonto '{metin(r['iskonto'])}'. Satır atlandı."))
                continue
            hazir.append({"kalem": metin(r["kalem"]), "tutar": v, "pb": pb_duzelt(r["pb"]), "iskonto": i, "kullanilabilir": yuvarla(v * (1 - i / 100))})
    return akislar, hazir, tatiller, uy


def projeksiyon(akislar: list[dict], acilis: Decimal, gunler: list[date], kesinti: Decimal = Decimal(0), artis: Decimal = Decimal(0)) -> list[dict]:
    """Gün gün giriş, çıkış, net, kümülatif ve gün sonu likidite. Kesinti / artış yalnız 'Beklenen' akışlara uygulanır."""
    satirlar, kum = [], Decimal(0)
    for g in gunler:
        gi = sum((a["tutar"] * (1 if a["kesin"] else 1 - kesinti / 100) for a in akislar if a["tarih"] == g and a["giris"]), Decimal(0))
        ci = sum((a["tutar"] * (1 if a["kesin"] else 1 + artis / 100) for a in akislar if a["tarih"] == g and not a["giris"]), Decimal(0))
        kum += gi - ci
        satirlar.append({"gun": g, "giris": yuvarla(gi), "cikis": yuvarla(ci), "net": yuvarla(gi - ci), "kumulatif": yuvarla(kum), "likidite": yuvarla(acilis + kum)})
    return satirlar


def ozetle(satirlar: list[dict], rapor: date) -> dict:
    en_dusuk = min(satirlar, key=lambda s: s["likidite"])
    ilk_acik = next((s for s in satirlar if s["likidite"] < 0), None)
    return {"en_dusuk": en_dusuk["likidite"], "en_dusuk_gun": en_dusuk["gun"], "ilk_acik": ilk_acik["gun"] if ilk_acik else None,
            "dayanma": (ilk_acik["gun"] - rapor).days if ilk_acik else None, "acik_gun": sum(s["likidite"] < 0 for s in satirlar)}


def analiz(akislar: list[dict], hazir: list[dict], rapor: date, gun_sayisi: int = 30, tatiller: set[date] = frozenset(), kesinti: Decimal = Decimal(50),
           artis: Decimal = Decimal(20), buyuk_esik: Decimal = Decimal(20)) -> dict:
    uy, son = [], rapor + timedelta(days=gun_sayisi - 1)
    gunler = [rapor + timedelta(days=i) for i in range(gun_sayisi) if is_gunu_mu(rapor + timedelta(days=i), set(tatiller))]
    if not gunler:
        raise ValueError("Rapor döneminde iş günü yok")
    disinda = sum(a["tarih"] > son for a in akislar)
    if disinda:
        uy.append(uyari("Bilgi", "Dönem dışı akış", f"{disinda} akış", f"{son:%d.%m.%Y} tarihinden sonraki akışlar rapora alınmadı (--gun ile dönemi uzatabilirsiniz)."))
    birimler = sorted({a["pb"] for a in akislar} | {h["pb"] for h in hazir}, key=lambda p: (p != "TL", p))
    sonuc = {}
    for pb in birimler:
        ak = [a for a in akislar if a["pb"] == pb and a["tarih"] <= son]
        acilis = sum((h["kullanilabilir"] for h in hazir if h["pb"] == pb), Decimal(0))
        baz, stres = projeksiyon(ak, acilis, gunler), projeksiyon(ak, acilis, gunler, kesinti, artis)
        haftalar = []
        for s, st in zip(baz, stres):
            hb = s["gun"] - timedelta(days=s["gun"].weekday())
            if not haftalar or haftalar[-1]["baslangic"] != hb:
                haftalar.append({"baslangic": hb, "ilk": s["gun"], "giris": Decimal(0), "cikis": Decimal(0)})
            h = haftalar[-1]
            h["giris"] += s["giris"]
            h["cikis"] += s["cikis"]
            h.update(son=s["gun"], net=h["giris"] - h["cikis"], likidite=s["likidite"], stres=st["likidite"])
            h["en_dusuk"] = min(h.get("en_dusuk", s["likidite"]), s["likidite"])
        ob, os_ = ozetle(baz, rapor), ozetle(stres, rapor)
        if not any(h["pb"] == pb for h in hazir):
            uy.append(uyari("Orta", "Hazır değer girilmemiş", pb, "Bu para biriminde açılış likiditesi 0 alındı; tablo yalnız akışların net etkisini gösterir."))
        if ob["ilk_acik"]:
            uy.append(uyari("Yüksek", "Likidite açığı", pb, f"Baz senaryoda {ob['ilk_acik']:%d.%m.%Y} günü likidite eksiye düşüyor; en düşük {bin_(ob['en_dusuk'])} "
                                                           f"({ob['en_dusuk_gun']:%d.%m.%Y}). Fonlama veya vade uzatımı planlayın."))
        elif os_["ilk_acik"]:
            uy.append(uyari("Orta", "Stres senaryosunda açık", pb, f"Beklenen girişler %{kesinti:f} az, beklenen çıkışlar %{artis:f} fazla gerçekleşirse "
                                                                  f"{os_['ilk_acik']:%d.%m.%Y} günü likidite eksiye düşüyor (en düşük {bin_(os_['en_dusuk'])})."))
        toplam_cikis = sum((a["tutar"] for a in ak if not a["giris"]), Decimal(0))
        toplam_giris = sum((a["tutar"] for a in ak if a["giris"]), Decimal(0))
        buyuk = [a | {"pay": yuvarla(a["tutar"] / (toplam_giris if a["giris"] else toplam_cikis) * 100, "0.1")} for a in ak
                 if (toplam_giris if a["giris"] else toplam_cikis) and a["tutar"] / (toplam_giris if a["giris"] else toplam_cikis) * 100 >= buyuk_esik]
        sonuc[pb] = {"acilis": acilis, "baz": baz, "stres": stres, "haftalar": haftalar, "ozet": ob, "ozet_stres": os_, "buyuk": sorted(buyuk, key=lambda a: -a["tutar"]),
                     "giris": toplam_giris, "cikis": toplam_cikis, "beklenen_giris": sum((a["tutar"] for a in ak if a["giris"] and not a["kesin"]), Decimal(0))}
    return {"pb": sonuc, "uyarilar": uy, "rapor": rapor, "son": son, "hazir": hazir, "kesinti": kesinti, "artis": artis}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

GUN_ADI = ["Pzt", "Sal", "Çar", "Per", "Cum", "Cmt", "Paz"]


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    oz = wb.active
    oz.title = "Özet"
    oz.append([f"Rapor tarihi {r['rapor']:%d.%m.%Y} · dönem sonu {r['son']:%d.%m.%Y} · stres: beklenen girişler −%{r['kesinti']:f}, beklenen çıkışlar +%{r['artis']:f}"])
    oz["A1"].font = Font(bold=True)
    _baslik(oz, ["Para Birimi", "Kullanılabilir Hazır Değer", "Toplam Giriş", "Beklenen Giriş Payı %", "Toplam Çıkış", "Dönem Sonu Likidite", "En Düşük Likidite", "En Düşük Gün",
                 "İlk Açık Günü", "Dayanma Süresi (gün)", "Stres: En Düşük", "Stres: İlk Açık Günü", "Stres: Dayanma (gün)"], (10, 20, 20, 12, 20, 20, 20, 12, 12, 11, 20, 12, 11), "B3")
    for pb, d in r["pb"].items():
        o, s = d["ozet"], d["ozet_stres"]
        _satir(oz, [pb, d["acilis"], d["giris"], yuvarla(d["beklenen_giris"] / d["giris"] * 100, "0.1") if d["giris"] else None, d["cikis"], d["baz"][-1]["likidite"], o["en_dusuk"],
                    o["en_dusuk_gun"], o["ilk_acik"], o["dayanma"] if o["dayanma"] is not None else "Dönem boyunca yeterli", s["en_dusuk"], s["ilk_acik"],
                    s["dayanma"] if s["dayanma"] is not None else "Dönem boyunca yeterli"], para_s=(2, 3, 5, 6, 7, 11), tarih_s=(8, 9, 12),
               dolgu={7: KIRMIZI if o["en_dusuk"] < 0 else YESIL, 11: KIRMIZI if s["en_dusuk"] < 0 else YESIL})
    for pb, d in r["pb"].items():
        gn = wb.create_sheet(f"Günlük {pb}")
        _baslik(gn, ["Tarih", "Gün", "Giriş", "Çıkış", "Net Akış", "Kümülatif Net", "Gün Sonu Likidite", "Stres: Gün Sonu Likidite"], (12, 6, 20, 20, 20, 20, 20, 22))
        _satir(gn, ["Açılış", "", "", "", "", "", d["acilis"], d["acilis"]], para_s=(7, 8))
        for b, s in zip(d["baz"], d["stres"]):
            _satir(gn, [b["gun"], GUN_ADI[b["gun"].weekday()], b["giris"], b["cikis"], b["net"], b["kumulatif"], b["likidite"], s["likidite"]], para_s=(3, 4, 5, 6, 7, 8), tarih_s=(1,),
                   dolgu={5: KIRMIZI if b["net"] < 0 else None, 7: KIRMIZI if b["likidite"] < 0 else None, 8: KIRMIZI if s["likidite"] < 0 else None})
        hf = wb.create_sheet(f"Haftalık {pb}")
        _baslik(hf, ["Hafta", "Giriş", "Çıkış", "Net Akış", "Hafta Sonu Likidite", "Hafta İçi En Düşük", "Stres: Hafta Sonu Likidite"], (26, 20, 20, 20, 20, 20, 22))
        for h in d["haftalar"]:
            _satir(hf, [f"{h['ilk']:%d.%m.%Y} – {h['son']:%d.%m.%Y}", h["giris"], h["cikis"], h["net"], h["likidite"], h["en_dusuk"], h["stres"]], para_s=(2, 3, 4, 5, 6, 7),
                   dolgu={4: KIRMIZI if h["net"] < 0 else None, 6: KIRMIZI if h["en_dusuk"] < 0 else None, 7: KIRMIZI if h["stres"] < 0 else None})
    by = wb.create_sheet("Büyük Akışlar")
    _baslik(by, ["Para Birimi", "Tarih", "Yön", "Kalem", "Karşı Taraf", "Tutar", "Dönem Toplamındaki Pay %", "Kesinlik"], (10, 12, 8, 40, 26, 20, 14, 10))
    for pb, d in r["pb"].items():
        for a in d["buyuk"]:
            _satir(by, [pb, a["tarih"], "Giriş" if a["giris"] else "Çıkış", a["kalem"], a["karsi"], a["tutar"], a["pay"], "Kesin" if a["kesin"] else "Beklenen"], para_s=(6,), tarih_s=(2,),
                   dolgu={8: None if a["kesin"] else SARI})
    hz = wb.create_sheet("Hazır Değerler")
    _baslik(hz, ["Kalem", "Para Birimi", "Tutar", "İskonto %", "Kullanılabilir"], (40, 10, 20, 10, 20))
    for h in r["hazir"]:
        _satir(hz, [h["kalem"], h["pb"], h["tutar"], h["iskonto"], h["kullanilabilir"]], para_s=(3, 5))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(a_yol: Path, cikti: Path, rapor: date, h_yol: Path | None = None, t_yol: Path | None = None, gun: int = 30, kesinti: Decimal = Decimal(50),
             artis: Decimal = Decimal(20)) -> dict:
    akislar, hazir, tatiller, uy = oku(a_yol, h_yol, rapor, t_yol)
    if not akislar and not hazir:
        raise ValueError(f"{a_yol.name}: okunabilir akış yok")
    r = analiz(akislar, hazir, rapor, gun, tatiller, kesinti, artis)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Vadesi gelen nakit akışlarından günlük ve haftalık likidite açığı / fazlası tablosu hazırlar.")
    p.add_argument("--akislar", type=Path, default=ORNEK / "akislar.csv", help="Tarih, Kalem, Yön (Giriş / Çıkış), Tutar, Para Birimi, Kesinlik (Kesin / Beklenen), Karşı Taraf")
    p.add_argument("--hazir", type=Path, help="Hazır değerler: Kalem, Tutar, Para Birimi, İskonto %%")
    p.add_argument("--tatiller", type=Path, help="Her satırda bir resmî tatil tarihi")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 05.10.2026)")
    p.add_argument("--gun", type=int, default=30, help="Rapor dönemi, takvim günü (varsayılan 30)")
    p.add_argument("--giris-kesinti", type=Decimal, default=Decimal(50), help="Stres: beklenen girişlerdeki azalma %% (varsayılan 50)")
    p.add_argument("--cikis-artis", type=Decimal, default=Decimal(20), help="Stres: beklenen çıkışlardaki artış %% (varsayılan 20)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "likidite_raporu.xlsx")
    a = p.parse_args(argv)
    ornek = a.akislar == ORNEK / "akislar.csv"
    hazir = a.hazir or (ORNEK / "hazir_degerler.csv" if ornek else None)
    tat = a.tatiller or (ORNEK / "tatiller.csv" if ornek else None)
    for y in (a.akislar, hazir, tat):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    rapor = tarih(a.tarih) if a.tarih else (date(2026, 10, 5) if ornek else date.today())
    if not rapor or not (1 <= a.gun <= 370) or not (0 <= a.giris_kesinti <= 100) or a.cikis_artis < 0:
        print("[X] --tarih GG.AA.YYYY; --gun 1-370; --giris-kesinti 0-100; --cikis-artis 0 veya üzeri olmalı")
        return 2
    try:
        r = calistir(a.akislar, a.cikti, rapor, hazir, tat, a.gun, a.giris_kesinti, a.cikis_artis)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    b = bin_
    for pb, d in r["pb"].items():
        o, s = d["ozet"], d["ozet_stres"]
        print(f"[OK] {pb}: açılış {b(d['acilis'])} · en düşük {b(o['en_dusuk'])} ({o['en_dusuk_gun']:%d.%m}) · " +
              (f"AÇIK {o['ilk_acik']:%d.%m.%Y}" if o["ilk_acik"] else "açık yok") + " · stres: " + (f"açık {s['ilk_acik']:%d.%m.%Y}" if s["ilk_acik"] else "açık yok"))
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['kim']} · {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
