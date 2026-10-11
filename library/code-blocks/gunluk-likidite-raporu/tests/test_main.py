import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
RAPOR = date(2026, 10, 5)                  # Pazartesi


def akis(gun, tutar, giris=True, kesin=True, pb="TL", kalem="K"):
    return {"tarih": date(2026, 10, gun), "asil": date(2026, 10, gun), "kalem": kalem, "giris": giris, "tutar": D(str(tutar)), "pb": pb, "kesin": kesin, "karsi": ""}


def hazir(tutar, iskonto=0, pb="TL"):
    return {"kalem": "H", "tutar": D(str(tutar)), "pb": pb, "iskonto": D(iskonto), "kullanilabilir": main.yuvarla(D(str(tutar)) * (1 - D(iskonto) / 100))}


class HesapTesti(unittest.TestCase):
    def test_projeksiyon(self):
        gunler = [date(2026, 10, g) for g in (5, 6, 7)]
        p = main.projeksiyon([akis(5, 100), akis(5, 250, False), akis(6, 400), akis(7, 300, False)], D(200), gunler)
        self.assertEqual([(s["net"], s["kumulatif"], s["likidite"]) for s in p], [(D(-150), D(-150), D(50)), (D(400), D(250), D(450)), (D(-300), D(-50), D(150))])

    def test_stres_yalniz_beklenene(self):
        a = [akis(5, 1000, True, True), akis(5, 1000, True, False), akis(5, 500, False, True), akis(5, 500, False, False)]
        s = main.projeksiyon(a, D(0), [RAPOR], D(50), D(20))[0]
        self.assertEqual((s["giris"], s["cikis"], s["net"]), (D("1500.00"), D("1100.00"), D("400.00")))   # 1000 + 500 ; 500 + 600

    def test_acik_ve_dayanma(self):
        r = main.analiz([akis(6, 300, False), akis(8, 500, False), akis(12, 900)], [hazir(1000, 40)], RAPOR, 10, kesinti=D(0), artis=D(0))
        d = r["pb"]["TL"]
        self.assertEqual(d["acilis"], D("600.00"))                                   # %40 iskonto
        self.assertEqual((d["ozet"]["ilk_acik"], d["ozet"]["dayanma"], d["ozet"]["en_dusuk"], d["ozet"]["acik_gun"]), (date(2026, 10, 8), 3, D("-200.00"), 2))
        self.assertEqual(d["baz"][-1]["likidite"], D("700.00"))                      # 12 Ekim'de giriş
        self.assertEqual([u["tur"] for u in r["uyarilar"]], ["Likidite açığı"])
        self.assertEqual(len(d["baz"]), 8)                                           # 5-14 Ekim arası iş günleri

    def test_yalniz_streste_acik(self):
        r = main.analiz([akis(6, 800, False, False), akis(6, 500, True, False)], [hazir(400)], RAPOR, 5, kesinti=D(50), artis=D(25))
        d = r["pb"]["TL"]
        self.assertEqual((d["ozet"]["ilk_acik"], d["ozet_stres"]["en_dusuk"]), (None, D("-350.00")))     # 400 + 250 − 1000
        self.assertEqual([u["tur"] for u in r["uyarilar"]], ["Stres senaryosunda açık"])

    def test_haftalik_ve_para_birimi(self):
        a = [akis(5, 100), akis(9, 50, False), akis(12, 30), akis(13, 500, False, pb="USD")]
        r = main.analiz(a, [hazir(10), hazir(600, pb="USD")], RAPOR, 14, kesinti=D(0), artis=D(0))
        self.assertEqual(list(r["pb"]), ["TL", "USD"])
        h = r["pb"]["TL"]["haftalar"]
        self.assertEqual([(x["ilk"], x["son"], x["net"], x["likidite"]) for x in h],
                         [(date(2026, 10, 5), date(2026, 10, 9), D(50), D(60)), (date(2026, 10, 12), date(2026, 10, 16), D(30), D(90))])
        self.assertEqual(r["pb"]["USD"]["baz"][-1]["likidite"], D("100.00"))

    def test_buyuk_akis_ve_donem_disi(self):
        a = [akis(6, 900, False, kalem="Büyük"), akis(7, 100, False), akis(7, 50), akis(30, 999)]
        r = main.analiz(a, [], RAPOR, 10, kesinti=D(0), artis=D(0))
        self.assertEqual([(x["kalem"], x["pay"]) for x in r["pb"]["TL"]["buyuk"]], [("Büyük", D("90.0")), ("K", D("100.0"))])
        self.assertEqual(sorted(u["tur"] for u in r["uyarilar"]), ["Dönem dışı akış", "Hazır değer girilmemiş", "Likidite açığı"])


class DosyaTesti(unittest.TestCase):
    def test_okuma(self):
        with tempfile.TemporaryDirectory() as tmp:
            a, h, t = (Path(tmp) / x for x in ("a.csv", "h.csv", "t.csv"))
            a.write_text("Tarih;Kalem;Yön;Tutar;Para Birimi;Kesinlik\n10.10.2026;Cumartesi;Giriş;100;TL;Kesin\n29.10.2026;Tatil;Çıkış;50;TRY;Beklenen\n"
                         "01.10.2026;Geçmiş;Tahsilat;30;;\n06.10.2026;Bozuk;Transfer;10;TL;\n07.10.2026;Eksi;Giriş;-5;TL;\n", encoding="utf-8")
            h.write_text("Kalem;Tutar;Para Birimi;İskonto\nNakit;1.000,00;TL;\nTahvil;2.000,00;TL;25\nBozuk;500;TL;150\n", encoding="utf-8")
            t.write_text("Tarih\n29.10.2026\n", encoding="utf-8")
            ak, hz, tat, uy = main.oku(a, h, RAPOR, t)
            self.assertEqual([(x["kalem"], x["tarih"], x["giris"], x["kesin"], x["pb"]) for x in ak],
                             [("Cumartesi", date(2026, 10, 12), True, True, "TL"), ("Tatil", date(2026, 10, 30), False, False, "TL"), ("Geçmiş", RAPOR, True, True, "TL")])
            self.assertEqual([x["kullanilabilir"] for x in hz], [D("1000.00"), D("1500.00")])
            self.assertEqual(sorted(u["tur"] for u in uy), ["Akış satırı okunamadı", "Akış satırı okunamadı", "Hazır değer okunamadı", "Tarihi geçmiş akış", "İş günü olmayan tarih"])
            arg = ["--akislar", str(a), "--hazir", str(h), "--tatiller", str(t), "--tarih", "05.10.2026", "--cikti", str(Path(tmp) / "x.xlsx")]
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(arg + ["--gun", "0"]), 2)
            self.assertEqual(main.main(["--akislar", str(Path(tmp) / "yok.csv")]), 1)

    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "l.xlsx"
            r = main.calistir(ORNEK / "akislar.csv", cikti, RAPOR, ORNEK / "hazir_degerler.csv", ORNEK / "tatiller.csv")
            tl, usd = r["pb"]["TL"], r["pb"]["USD"]
            self.assertEqual((tl["acilis"], usd["acilis"]), (D("2660000000.00"), D("63500000.00")))
            self.assertEqual((tl["ozet"]["ilk_acik"], usd["ozet"]["ilk_acik"], usd["ozet"]["dayanma"]), (None, date(2026, 10, 16), 11))
            self.assertIsNotNone(tl["ozet_stres"]["ilk_acik"])
            self.assertNotIn(date(2026, 10, 29), [s["gun"] for s in tl["baz"]])       # tatil günü tabloda yok
            self.assertEqual(tl["baz"][-1]["likidite"], tl["acilis"] + tl["giris"] - tl["cikis"])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Günlük TL", "Haftalık TL", "Günlük USD", "Haftalık USD", "Büyük Akışlar", "Hazır Değerler", "Uyarılar"])
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "o.xlsx")]), 0)


if __name__ == "__main__":
    unittest.main()
