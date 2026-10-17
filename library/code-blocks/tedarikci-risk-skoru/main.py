"""
Tedarikçi Risk Skoru — Gerisi Bende kod bloğu
Satın Alma › Satın Alma Müdürü

Tek kaynak bağımlılığı, finansal durum, gecikme ve kalite verisinden tedarikçi başına 0–100 arası risk skoru üretir.
  Alt puanlar (0 = risksiz, 100 = en riskli)
    Tek kaynak   = tek kaynak kalem sayısı ÷ kalem sayısı × 100
    Finansal     = İyi 0 · Orta 50 · Zayıf 100  (değerlendirmeyi siz yaparsınız)
    Gecikme      = gecikmeli teslimat % ÷ --gecikme-tavan (20), en çok 100
    Kalite       = ret % ÷ --ret-tavan (5), en çok 100
  Risk skoru     = ağırlıklı ortalama; ağırlıklar --agirlik (tek 30, finans 25, gecikme 25, kalite 20).
                   Verisi olmayan ölçüt hesaptan çıkarılır, kalan ağırlıklar yeniden oranlanır
  Sınıf          = skor ≥ --yuksek (60) Yüksek · ≥ --orta (30) Orta · altı Düşük · 2'den az ölçüt varsa Değerlendirilemedi
  Harcama payı   = tedarikçiden yıllık alım ÷ toplam alım
Skor bir önceliklendirme aracıdır; ağırlıklar ve eşikler yönetim tercihidir, nesnel bir kredi notu değildir.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
SUTUN = {"ad": ("tedarikci", "tedarikci adi", "firma", "unvan"), "alim": ("yillik alim tutari", "alim tutari", "yillik alim", "ciro", "harcama"), "kalem": ("kalem sayisi", "malzeme sayisi", "kalem"),
         "tek": ("tek kaynak kalem", "tek kaynak kalem sayisi", "tek kaynak"), "gecikme": ("gecikmeli teslimat", "gecikme", "gecikme orani", "gec teslimat"), "ret": ("ret", "ret orani", "iade", "iade orani", "kalite ret"),
         "finans": ("finansal durum", "finansal degerlendirme", "mali durum")}
FINANS = {"iyi": 0, "guclu": 0, "saglam": 0, "orta": 50, "zayif": 100, "kotu": 100, "riskli": 100}
FINANS_AD = {0: "İyi", 50: "Orta", 100: "Zayıf"}
OLCUT = {"tek": "Tek kaynak", "finans": "Finansal", "gecikme": "Gecikme", "kalite": "Kalite"}
VARSAYILAN = {"tek": Decimal(30), "finans": Decimal(25), "gecikme": Decimal(25), "kalite": Decimal(20)}
S_RENK = {"Yüksek": KIRMIZI, "Orta": SARI, "Düşük": YESIL}
Z, YUZ = Decimal(0), Decimal(100)


def olcu(x) -> Decimal | None:
    """Yüzde ve adetler: virgül veya tek nokta ondalık ayırıcıdır (1.5 = 1,5)."""
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x or ""))
    s = s.replace(".", "").replace(",", ".") if "," in s else s
    try:
        return Decimal(s) if s else None
    except InvalidOperation:
        return None


def agirlik_coz(s: str) -> dict | None:
    a = dict(VARSAYILAN)
    for parca in (x for x in re.split(r"[;\s]+", s.strip()) if x):
        m = re.fullmatch(r"([a-zçğıöşü]+)=(.+)", parca.lower())
        v = olcu(m.group(2)) if m else None
        if not m or katla(m.group(1)) not in a or v is None or v < 0:
            return None
        a[katla(m.group(1))] = v
    return a if sum(a.values()) > 0 else None


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("ad",)):
        ad = metin(r["ad"])
        if not ad:
            continue
        kalem, tek, gec, ret, alim, f = olcu(r["kalem"]), olcu(r["tek"]), olcu(r["gecikme"]), olcu(r["ret"]), para(r["alim"]), katla(r["finans"])
        hata = [n for n, v, ust in (("Kalem Sayısı", kalem, None), ("Tek Kaynak Kalem", tek, kalem), ("Gecikmeli Teslimat %", gec, YUZ), ("Ret %", ret, YUZ)) if v is not None and (v < 0 or (ust is not None and v > ust))]
        if f and f not in FINANS:
            hata.append("Finansal Durum (İyi / Orta / Zayıf)")
        if hata:
            uy.append(uyari("Orta", "Değer geçersiz", ad, f"Şu alanlar geçersiz olduğu için yok sayıldı: {', '.join(hata)}."))
        gecerli = lambda n, v: None if n in hata else v   # noqa: E731
        l.append({"ad": ad, "alim": alim if alim and alim > 0 else None, "kalem": gecerli("Kalem Sayısı", kalem), "tek": gecerli("Tek Kaynak Kalem", tek), "gecikme": gecerli("Gecikmeli Teslimat %", gec), "ret": gecerli("Ret %", ret),
                  "finans": f if f in FINANS else ""})
    return l, uy


def skorla(tedarikciler: list[dict], agirlik: dict | None = None, gecikme_tavan: Decimal = Decimal(20), ret_tavan: Decimal = Decimal(5), yuksek: Decimal = Decimal(60), orta: Decimal = Decimal(30), pay_esik: Decimal = Decimal(10),
           yogunlasma: Decimal = Decimal(30)) -> dict:
    if not tedarikciler:
        raise ValueError("Skorlanacak tedarikçi yok.")
    agirlik, uy = agirlik or VARSAYILAN, []
    toplam = sum((t["alim"] for t in tedarikciler if t["alim"]), Z)
    for t in tedarikciler:
        p = {"tek": t["tek"] * YUZ / t["kalem"] if t["tek"] is not None and t["kalem"] else None, "finans": Decimal(FINANS[t["finans"]]) if t["finans"] else None,
             "gecikme": min(t["gecikme"] / gecikme_tavan, Decimal(1)) * YUZ if t["gecikme"] is not None else None, "kalite": min(t["ret"] / ret_tavan, Decimal(1)) * YUZ if t["ret"] is not None else None}
        var = [k for k in OLCUT if p[k] is not None and agirlik[k] > 0]
        t["puan"] = {k: yuvarla(v) if v is not None else None for k, v in p.items()}
        t["eksik"] = [OLCUT[k] for k in OLCUT if p[k] is None]
        t["pay"] = yuvarla(t["alim"] * YUZ / toplam) if t["alim"] and toplam else None
        if len(var) < 2:
            t["skor"], t["sinif"] = None, "Değerlendirilemedi"
            uy.append(uyari("Orta", "Tedarikçi değerlendirilemedi", t["ad"], f"Yalnız {len(var)} ölçütün verisi var; skor için en az 2 ölçüt gerekir. Eksik: {', '.join(t['eksik'])}."))
            continue
        t["skor"] = yuvarla(sum((agirlik[k] * p[k] for k in var), Z) / sum((agirlik[k] for k in var), Z))
        t["sinif"] = "Yüksek" if t["skor"] >= yuksek else "Orta" if t["skor"] >= orta else "Düşük"
        if t["eksik"]:
            uy.append(uyari("Bilgi", "Eksik veri", t["ad"], f"{', '.join(t['eksik'])} ölçütünün verisi yok; skor kalan ölçütlerle hesaplandı."))
        if t["sinif"] == "Yüksek":
            kritik = t["pay"] is not None and t["pay"] >= pay_esik
            uy.append(uyari("Yüksek" if kritik else "Orta", "Kritik tedarikçi" if kritik else "Yüksek riskli tedarikçi", t["ad"],
                            f"Risk skoru {t['skor']}" + (f"; toplam alımın %{t['pay']} kadarı bu tedarikçiden." if t["pay"] is not None else ".") + " Alternatif kaynak ve iyileştirme planı hazırlayın."))
        if t["tek"] and t["finans"] and FINANS[t["finans"]] == 100:
            uy.append(uyari("Yüksek", "Tek kaynak ve finansal durumu zayıf", t["ad"], f"{t['tek']} kalemde tek kaynak ve finansal durumu zayıf değerlendirilmiş. Tedarik kesintisi riskine karşı ikinci kaynak geliştirin."))
        if t["pay"] is not None and t["pay"] >= yogunlasma and len(tedarikciler) > 1:
            uy.append(uyari("Bilgi", "Harcama yoğunlaşması", t["ad"], f"Toplam alımın %{t['pay']} kadarı bu tedarikçiden."))
    sira = {"Yüksek": 0, "Orta": 1, "Düşük": 2, "Değerlendirilemedi": 3}
    tedarikciler.sort(key=lambda t: (sira[t["sinif"]], -(t["skor"] or Z), -(t["alim"] or Z), katla(t["ad"])))
    return {"tedarikciler": tedarikciler, "say": {s: sum(t["sinif"] == s for t in tedarikciler) for s in sira}, "toplam": toplam, "yuksek_alim": sum((t["alim"] for t in tedarikciler if t["sinif"] == "Yüksek" and t["alim"]), Z),
            "agirlik": agirlik, "param": {"gecikme_tavan": gecikme_tavan, "ret_tavan": ret_tavan, "yuksek": yuksek, "orta": orta}, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Risk Skoru"
    _baslik(ws, ["Sınıf", "Tedarikçi", "Risk Skoru", "Tek Kaynak Puanı", "Finansal Puan", "Gecikme Puanı", "Kalite Puanı", "Yıllık Alım", "Harcama Payı %", "Kalem", "Tek Kaynak Kalem", "Gecikmeli Teslimat %", "Ret %", "Finansal Durum",
                 "Eksik Veri"], (18, 30, 10, 10, 10, 10, 10, 16, 10, 8, 10, 11, 8, 12, 26))
    for t in r["tedarikciler"]:
        _satir(ws, [t["sinif"], t["ad"], t["skor"], t["puan"]["tek"], t["puan"]["finans"], t["puan"]["gecikme"], t["puan"]["kalite"], t["alim"], t["pay"], t["kalem"], t["tek"], t["gecikme"], t["ret"], FINANS_AD[FINANS[t["finans"]]] if t["finans"] else "",
                    ", ".join(t["eksik"])], para_s=(3, 4, 5, 6, 7, 8, 9), dolgu={1: S_RENK.get(t["sinif"]), 3: S_RENK.get(t["sinif"])})
    ws.append([])
    a, p = r["agirlik"], r["param"]
    ws.append([f"Ağırlıklar: tek kaynak {a['tek']}, finansal {a['finans']}, gecikme {a['gecikme']}, kalite {a['kalite']} · gecikme tavanı %{p['gecikme_tavan']}, ret tavanı %{p['ret_tavan']} · Yüksek ≥ {p['yuksek']}, Orta ≥ {p['orta']}. "
               "Puanlar 0 (risksiz) – 100 (en riskli)."])
    ws.append([f"Toplam alım {r['toplam']:,.2f}; yüksek riskli tedarikçilerden alım {r['yuksek_alim']:,.2f}" + (f" (%{yuvarla(r['yuksek_alim'] * YUZ / r['toplam'])})." if r["toplam"] else ".")])
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, **secenek) -> dict:
    tedarikciler, uy = oku(yol)
    r = skorla(tedarikciler, **secenek)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Tedarikçi başına 0–100 arası risk skoru üretir.")
    p.add_argument("--tedarikciler", type=Path, help="Tedarikçi, Yıllık Alım Tutarı, Kalem Sayısı, Tek Kaynak Kalem, Gecikmeli Teslimat %%, Ret %%, Finansal Durum")
    p.add_argument("--agirlik", default="", help='Ölçüt ağırlıkları, örn. "tek=40;finans=20;gecikme=20;kalite=20" (varsayılan 30;25;25;20)')
    p.add_argument("--gecikme-tavan", default="20", help="Gecikmeli teslimat bu yüzdeye ulaşınca puan 100 olur (varsayılan 20)")
    p.add_argument("--ret-tavan", default="5", help="Ret oranı bu yüzdeye ulaşınca puan 100 olur (varsayılan 5)")
    p.add_argument("--yuksek", default="60", help="Skor bu değer ve üstündeyse Yüksek risk (varsayılan 60)")
    p.add_argument("--orta", default="30", help="Skor bu değer ve üstündeyse Orta risk (varsayılan 30)")
    p.add_argument("--pay", default="10", help="Yüksek riskli tedarikçinin harcama payı bu yüzdeye ulaşıyorsa 'kritik' sayılır (varsayılan 10)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "tedarikci_risk_skoru.xlsx")
    a = p.parse_args(argv)
    yol = a.tedarikciler or ORNEK / "tedarikciler.csv"
    agirlik, gt, rt, yuksek, orta, pay = agirlik_coz(a.agirlik), olcu(a.gecikme_tavan), olcu(a.ret_tavan), olcu(a.yuksek), olcu(a.orta), olcu(a.pay)
    if agirlik is None or None in (gt, rt, yuksek, orta, pay) or gt <= 0 or rt <= 0 or not 0 < orta < yuksek <= 100 or not 0 <= pay <= 100:
        print('[X] --agirlik "tek=30;finans=25;gecikme=25;kalite=20" biçiminde; tavanlar pozitif; --orta sıfırdan büyük ve --yuksek değerinden küçük; --pay 0–100 arasında olmalı')
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, agirlik=agirlik, gecikme_tavan=gt, ret_tavan=rt, yuksek=yuksek, orta=orta, pay_esik=pay)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    s = r["say"]
    print(f"[OK] {len(r['tedarikciler'])} tedarikçi · Yüksek {s['Yüksek']}, Orta {s['Orta']}, Düşük {s['Düşük']}" + (f", değerlendirilemeyen {s['Değerlendirilemedi']}" if s["Değerlendirilemedi"] else "")
          + f" · yüksek riskli tedarikçilerden alım {r['yuksek_alim']:,.2f}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
