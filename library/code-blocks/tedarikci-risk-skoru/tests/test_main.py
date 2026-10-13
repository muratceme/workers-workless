import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def t(ad, kalem=None, tek=None, gecikme=None, ret=None, finans="", alim=None):
    d = lambda x: None if x is None else D(str(x))   # noqa: E731
    return {"ad": ad, "alim": d(alim), "kalem": d(kalem), "tek": d(tek), "gecikme": d(gecikme), "ret": d(ret), "finans": finans}


def tek(*a, **kw):
    secenek = {k: kw.pop(k) for k in list(kw) if k in ("agirlik", "gecikme_tavan", "ret_tavan", "yuksek", "orta")}
    x = main.skorla([t("T", *a, **kw)], **secenek)["tedarikciler"][0]
    return x["skor"], x["sinif"], x["puan"]


class SkorTesti(unittest.TestCase):
    def test_puanlar_ve_skor(self):
        self.assertEqual(tek(8, 4, 15, 4, "zayif"), (D("74.75"), "Yüksek", {"tek": D("50.00"), "finans": D("100.00"), "gecikme": D("75.00"), "kalite": D("80.00")}))
        self.assertEqual(tek(3, 1, 25, 7.5, "orta"), (D("67.50"), "Yüksek", {"tek": D("33.33"), "finans": D("50.00"), "gecikme": D("100.00"), "kalite": D("100.00")}))      # (30×33,33 + 25×50 + 25×100 + 20×100) ÷ 100
        self.assertEqual(tek(10, 0, 0, 0, "iyi"), (D("0.00"), "Düşük", {"tek": D("0.00"), "finans": D("0.00"), "gecikme": D("0.00"), "kalite": D("0.00")}))
        self.assertEqual([tek(10, 10, 20, s, "zayif")[:2] for s in (0, 2.5)], [(D("80.00"), "Yüksek"), (D("90.00"), "Yüksek")])
        self.assertEqual([tek(10, k, 0, 0, "iyi")[:2] for k in (10, 9.9)], [(D("30.00"), "Orta"), (D("29.70"), "Düşük")])                                              # Orta eşiği 30
        self.assertEqual([tek(10, 10, g, 0, "zayif")[:2] for g in (4, 3.9)], [(D("60.00"), "Yüksek"), (D("59.88"), "Orta")])                                           # Yüksek eşiği 60

    def test_eksik_veri_ve_ayarlar(self):
        self.assertEqual(tek(3, 0, 12, None, "iyi")[:2], (D("18.75"), "Düşük"))                           # kalite yok: (25×60) ÷ 80
        self.assertEqual(tek(None, None, 10, 5)[:2], (D("72.22"), "Yüksek"))                              # yalnız gecikme 50 ve kalite 100: (25×50 + 20×100) ÷ 45
        self.assertEqual(tek(0, 0, 10, 5)[2]["tek"], None)                                                # kalem sayısı 0 → tek kaynak puanı hesaplanamaz
        self.assertEqual((tek(6, 0)[:2], tek()[:2]), ((None, "Değerlendirilemedi"), (None, "Değerlendirilemedi")))
        self.assertEqual(tek(8, 4, 15, 4, "zayif", agirlik={"tek": D(1), "finans": D(0), "gecikme": D(0), "kalite": D(1)})[:2], (D("65.00"), "Yüksek"))
        self.assertEqual(tek(8, 4, 15, 4, "zayif", gecikme_tavan=D(30), ret_tavan=D(8))[:2], (D("62.50"), "Yüksek"))      # gecikme 50, kalite 50
        self.assertEqual(tek(8, 4, 15, 4, "zayif", yuksek=D(75), orta=D("74.75"))[1], "Orta")
        r = main.skorla([t("Eksik", 3, 0, 12, None, "iyi"), t("Tek", 6, 0)])
        self.assertEqual(([(x["ad"], x["eksik"]) for x in r["tedarikciler"]], sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"])),
                         ([("Eksik", ["Kalite"]), ("Tek", ["Finansal", "Gecikme", "Kalite"])], [("Bilgi", "Eksik veri", "Eksik"), ("Orta", "Tedarikçi değerlendirilemedi", "Tek")]))
        with self.assertRaises(ValueError):
            main.skorla([])

    def test_pay_ve_uyarilar(self):
        l = [t("Kritik", 5, 5, 20, 5, "zayif", 100), t("Küçük", 5, 5, 20, 5, "zayif", 99), t("Büyük", 10, 0, 0, 0, "iyi", 600), t("Orta", 10, 5, 10, 1, "orta", 200), t("Paysız", 5, 5, 20, 5, "orta"), t("Sınır", 10, 0, 0, 0, "iyi", 1)]
        r = main.skorla(l)
        self.assertEqual([(x["ad"], x["sinif"], x["skor"], x["pay"]) for x in r["tedarikciler"]],
                         [("Kritik", "Yüksek", D("100.00"), D("10.00")), ("Küçük", "Yüksek", D("100.00"), D("9.90")), ("Paysız", "Yüksek", D("87.50"), None), ("Orta", "Orta", D("44.00"), D("20.00")),
                          ("Büyük", "Düşük", D("0.00"), D("60.00")), ("Sınır", "Düşük", D("0.00"), D("0.10"))])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Bilgi", "Harcama yoğunlaşması", "Büyük"), ("Orta", "Yüksek riskli tedarikçi", "Küçük"), ("Orta", "Yüksek riskli tedarikçi", "Paysız"), ("Yüksek", "Kritik tedarikçi", "Kritik"),
                          ("Yüksek", "Tek kaynak ve finansal durumu zayıf", "Kritik"), ("Yüksek", "Tek kaynak ve finansal durumu zayıf", "Küçük")])
        self.assertEqual((r["say"], r["toplam"], r["yuksek_alim"]), ({"Yüksek": 3, "Orta": 1, "Düşük": 2, "Değerlendirilemedi": 0}, D(1000), D(199)))


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "tedarikciler.csv", cikti)
            self.assertEqual([(x["ad"], x["sinif"], x["skor"], x["pay"]) for x in r["tedarikciler"]],
                             [("Kalıp G Ltd.", "Yüksek", D("100.00"), D("3.49")), ("Elektronik E Ltd.", "Yüksek", D("74.75"), D("15.81")), ("Kimya B A.Ş.", "Yüksek", D("73.50"), D("21.40")),
                              ("Demir Çelik A.Ş.", "Düşük", D("28.20"), D("43.02")), ("Lojistik F A.Ş.", "Düşük", D("18.75"), D("5.58")), ("Ambalaj A Ltd.", "Düşük", D("7.00"), D("9.53")),
                              ("Matbaa D Ltd.", "Değerlendirilemedi", None, D("1.16"))])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Yüksek", "Kritik tedarikçi", "Elektronik E Ltd."), ("Yüksek", "Kritik tedarikçi", "Kimya B A.Ş."), ("Yüksek", "Tek kaynak ve finansal durumu zayıf", "Elektronik E Ltd."),
                              ("Yüksek", "Tek kaynak ve finansal durumu zayıf", "Kalıp G Ltd."), ("Orta", "Tedarikçi değerlendirilemedi", "Matbaa D Ltd."), ("Orta", "Yüksek riskli tedarikçi", "Kalıp G Ltd."),
                              ("Bilgi", "Eksik veri", "Lojistik F A.Ş."), ("Bilgi", "Harcama yoğunlaşması", "Demir Çelik A.Ş.")])
            self.assertEqual((r["toplam"], r["yuksek_alim"]), (D(43000000), D(17500000)))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Risk Skoru", "Uyarılar"])

    def test_okuma_ve_cli(self):
        self.assertEqual((main.agirlik_coz(""), main.agirlik_coz("tek=40; kalite=10,5"), main.agirlik_coz("TEK=1 finans=0 gecikme=0 kalite=0")),
                         (main.VARSAYILAN, {"tek": D(40), "finans": D(25), "gecikme": D(25), "kalite": D("10.5")}, {"tek": D(1), "finans": D(0), "gecikme": D(0), "kalite": D(0)}))
        self.assertEqual([main.agirlik_coz(x) for x in ("fiyat=10", "tek=-1", "tek", "tek=x", "tek=0;finans=0;gecikme=0;kalite=0")], [None] * 5)
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "t.csv"
            y.write_text("Firma;Harcama;Malzeme Sayısı;Tek Kaynak;Gecikme Oranı;İade Oranı;Mali Durum\nA;1.250.000,50;10;3;%12,5;1.5;GÜÇLÜ\nB;;5;6;120;-1;belirsiz\nC;x;;;;;\n;5;;;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["ad"], x["alim"], x["kalem"], x["tek"], x["gecikme"], x["ret"], x["finans"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("A", D("1250000.50"), D(10), D(3), D("12.5"), D("1.5"), "guclu"), ("B", None, D(5), None, None, None, ""), ("C", None, None, None, None, None, "")], [("Değer geçersiz", "B")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--tedarikciler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--agirlik", "tek=40;finans=20;gecikme=20;kalite=20", "--gecikme-tavan", "15", "--ret-tavan", "2,5", "--yuksek", "50", "--orta", "20", "--pay", "5"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + x) for x in (["--agirlik", "x=1"], ["--gecikme-tavan", "0"], ["--ret-tavan", "x"], ["--yuksek", "30"], ["--orta", "0"], ["--yuksek", "101"], ["--pay", "101"])], [2] * 7)
            self.assertEqual(main.main(["--tedarikciler", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Kalem;Ret\n1;1\n", "Tedarikçi;Ret\n;1\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
