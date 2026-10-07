"""
Alacak Yaşlandırma — Workers / Workless kod bloğu
Finans › Finans Uzmanı

Cari hesap hareketlerinden (borç/alacak) tahsilatları en eski faturadan başlayarak kapatır (FIFO),
açık kalan faturaları vade gününe göre yaşlandırır: vadesi gelmemiş, 1-30, 31-60, 61-90, 91-180, 180+ gün.
Müşteri bazında ağırlıklı ortalama gecikme, kredi limiti aşımı ve riskli alacakları işaretler.
Hazır "açık kalemler" listesi (Kalan sütunu) verilirse doğrudan onu yaşlandırır. İnternete bağlanmaz.

Kullanım:
    python main.py                                           # örnek cari hareketlerle dener
    python main.py --girdi cari_hareketler.xlsx --tarih 30.09.2026
    python main.py --girdi cari_hareketler.xlsx --limitler kredi_limitleri.xlsx --vade 30
"""
from __future__ import annotations

import argparse
import csv
import sys
from collections import OrderedDict, defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
DILIMLER = [("Vadesi gelmemiş", None, 0), ("1-30 gün", 1, 30), ("31-60 gün", 31, 60), ("61-90 gün", 61, 90),
            ("91-180 gün", 91, 180), ("180+ gün", 181, None)]


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def para(x) -> Decimal:
    if x in (None, ""):
        return Decimal(0)
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("₺", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return Decimal(0)


def tarih(x) -> date | None:
    if x in (None, ""):
        return None
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(str(x).strip(), f).date()
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ayirici = csv.Sniffer().sniff(metin.splitlines()[0], delimiters=",;\t").delimiter
        satirlar = list(csv.reader(metin.splitlines(), delimiter=ayirici))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


ALANLAR = {
    "cari": ("cari", "cari unvan", "müşteri", "cari hesap", "cari adı", "unvan"),
    "cari_kod": ("cari kod", "cari kodu", "müşteri kodu", "hesap kodu"),
    "tarih": ("tarih", "belge tarihi", "fatura tarihi", "işlem tarihi"),
    "vade": ("vade", "vade tarihi", "son ödeme tarihi"),
    "belge": ("belge no", "fatura no", "evrak no", "fiş no"),
    "aciklama": ("açıklama",),
    "borc": ("borç", "borc", "tutar (borç)"),
    "alacak": ("alacak", "tutar (alacak)"),
    "kalan": ("kalan", "açık tutar", "bakiye", "kalan tutar"),
    "tutar": ("tutar", "fatura tutarı"),
}


@dataclass
class Kalem:
    cari: str
    belge: str
    tarih: date
    vade: date
    tutar: Decimal
    kalan: Decimal


def acik_kalemler(satirlar: list[list], varsayilan_vade: int) -> tuple[list[Kalem], dict, list[str]]:
    b = [kucuk(x) for x in satirlar[0]]
    k = {alan: next((i for i, x in enumerate(b) if x in adlar), None) for alan, adlar in ALANLAR.items()}
    if k["cari"] is None and k["cari_kod"] is None:
        raise SystemExit(f"'Cari' sütunu bulunamadı. Başlıklar: {satirlar[0]}")
    al = lambda r, a: r[k[a]] if k[a] is not None and k[a] < len(r) else None  # noqa: E731
    cari_adi = lambda r: str(al(r, "cari") or al(r, "cari_kod") or "").strip()  # noqa: E731
    uyarilar = []
    fazla_tahsilat: dict[str, Decimal] = defaultdict(Decimal)

    if k["kalan"] is not None:            # hazır açık kalem listesi
        kalemler = []
        for r in satirlar[1:]:
            t = tarih(al(r, "tarih"))
            kalan = para(al(r, "kalan"))
            if kalan == 0 or t is None:
                continue
            v = tarih(al(r, "vade")) or t + timedelta(days=varsayilan_vade)
            kalemler.append(Kalem(cari_adi(r), str(al(r, "belge") or ""), t, v, para(al(r, "tutar")) or kalan, kalan))
        return kalemler, fazla_tahsilat, uyarilar

    if k["borc"] is None or k["alacak"] is None or k["tarih"] is None:
        raise SystemExit("Hareket listesinde Tarih, Borç ve Alacak sütunları (veya açık kalem listesinde Kalan) gerekli.")
    hareketler: "OrderedDict[str, list]" = OrderedDict()
    for r in satirlar[1:]:
        t = tarih(al(r, "tarih"))
        if t is None:
            continue
        hareketler.setdefault(cari_adi(r), []).append((t, tarih(al(r, "vade")), str(al(r, "belge") or ""),
                                                       para(al(r, "borc")), para(al(r, "alacak"))))
    kalemler = []
    for cari, liste in hareketler.items():
        liste.sort(key=lambda x: x[0])
        acik: list[Kalem] = []
        alacak_havuzu = Decimal(0)
        for t, v, belge, borc, alacak in liste:
            if borc > 0:
                acik.append(Kalem(cari, belge, t, v or t + timedelta(days=varsayilan_vade), borc, borc))
            alacak_havuzu += alacak
            # FIFO: tahsilatlar en eski açık faturayı kapatır
            for kalem in acik:
                if alacak_havuzu <= 0:
                    break
                kapat = min(kalem.kalan, alacak_havuzu)
                kalem.kalan -= kapat
                alacak_havuzu -= kapat
            acik = [x for x in acik if x.kalan > 0]
        if alacak_havuzu > 0:
            fazla_tahsilat[cari] += alacak_havuzu
            uyarilar.append(f"{cari}: {alacak_havuzu} TL tahsilat açık faturayı aşıyor (avans / alacaklı bakiye)")
        kalemler += acik
    return kalemler, fazla_tahsilat, uyarilar


def dilim(gecikme: int) -> str:
    for ad, alt, ust in DILIMLER:
        if alt is None and gecikme <= 0:
            return ad
        if alt is not None and gecikme >= alt and (ust is None or gecikme <= ust):
            return ad
    return DILIMLER[-1][0]


def limitleri_oku(yol: Path | None) -> dict[str, Decimal]:
    if not yol:
        return {}
    satirlar = tablo_oku(yol)
    b = [kucuk(x) for x in satirlar[0]]
    i_c = next((i for i, x in enumerate(b) if x in ALANLAR["cari"] + ALANLAR["cari_kod"]), None)
    i_l = next((i for i, x in enumerate(b) if x in ("kredi limiti", "limit", "risk limiti")), None)
    if i_c is None or i_l is None:
        raise SystemExit(f"Limit dosyasında Cari ve Kredi Limiti sütunları gerekli. Başlıklar: {satirlar[0]}")
    return {str(r[i_c]).strip(): para(r[i_l]) for r in satirlar[1:]}


BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
PARA = "#,##0.00"


def calistir(girdi: Path, rapor_tarihi: date, cikti: Path, varsayilan_vade: int = 0, limit_yolu: Path | None = None) -> dict:
    kalemler, fazla, uyarilar = acik_kalemler(tablo_oku(girdi), varsayilan_vade)
    limitler = limitleri_oku(limit_yolu)
    musteri: "OrderedDict[str, dict]" = OrderedDict()
    for x in sorted(kalemler, key=lambda x: (x.cari, x.vade)):
        g = (rapor_tarihi - x.vade).days
        m = musteri.setdefault(x.cari, {d[0]: Decimal(0) for d in DILIMLER} | {"toplam": Decimal(0), "agirlik": Decimal(0), "vadesi_gecmis": Decimal(0)})
        m[dilim(g)] += x.kalan
        m["toplam"] += x.kalan
        if g > 0:
            m["vadesi_gecmis"] += x.kalan
            m["agirlik"] += x.kalan * g

    wb = Workbook()
    ws = wb.active
    ws.title = "Müşteri Yaşlandırma"
    basliklar = ["Cari", "Toplam Açık"] + [d[0] for d in DILIMLER] + ["Vadesi Geçmiş", "Ağırlıklı Ort. Gecikme (gün)",
                                                                       "Avans / Alacaklı", "Kredi Limiti", "Limit Durumu", "Risk"]
    ws.append(basliklar)
    for h in ws[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    genel = defaultdict(Decimal)
    for cari, m in sorted(musteri.items(), key=lambda i: -i[1]["vadesi_gecmis"]):
        ort = float(m["agirlik"] / m["vadesi_gecmis"]) if m["vadesi_gecmis"] else 0.0
        limit = limitler.get(cari)
        limit_durum = "" if limit is None else ("AŞILDI" if m["toplam"] > limit else f"%{m['toplam'] / limit * 100:.0f} kullanım" if limit else "")
        risk = []
        if m["91-180 gün"] + m["180+ gün"] > 0:
            risk.append("90 günü aşan alacak")
        if m["180+ gün"] > 0:
            risk.append("şüpheli alacak değerlendirmesi")
        if limit_durum == "AŞILDI":
            risk.append("limit aşımı")
        ws.append([cari, m["toplam"]] + [m[d[0]] for d in DILIMLER] + [m["vadesi_gecmis"], round(ort, 1),
                                                                      fazla.get(cari, Decimal(0)), limit, limit_durum, "; ".join(risk)])
        if risk:
            ws.cell(ws.max_row, len(basliklar)).fill = KIRMIZI
        for d in [d[0] for d in DILIMLER] + ["toplam", "vadesi_gecmis"]:
            genel[d] += m[d]
    ws.append(["TOPLAM", genel["toplam"]] + [genel[d[0]] for d in DILIMLER] + [genel["vadesi_gecmis"]])
    for h in ws[ws.max_row]:
        h.font = Font(bold=True)
    for r in ws.iter_rows(min_row=2):
        for h in r:
            if isinstance(h.value, Decimal):
                h.value = float(h.value)
                h.number_format = PARA
    for j, b in enumerate(basliklar, 1):
        ws.column_dimensions[get_column_letter(j)].width = 30 if j == 1 else 16
    ws.freeze_panes = "B2"

    if musteri:
        g = BarChart()
        g.type, g.grouping, g.overlap = "bar", "stacked", 100
        g.title, g.height, g.width = "Müşteri bazında yaşlandırma (ilk 15)", 10, 22
        n = min(15, len(musteri))
        g.add_data(Reference(ws, min_col=3, max_col=2 + len(DILIMLER), min_row=1, max_row=1 + n), titles_from_data=True)
        g.set_categories(Reference(ws, min_col=1, min_row=2, max_row=1 + n))
        ws.add_chart(g, f"B{ws.max_row + 3}")

    d = wb.create_sheet("Açık Faturalar")
    d.append(["Cari", "Belge No", "Belge Tarihi", "Vade", "Tutar", "Açık Kalan", "Gecikme (gün)", "Dilim"])
    for h in d[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for x in sorted(kalemler, key=lambda x: (x.vade, x.cari)):
        g = (rapor_tarihi - x.vade).days
        d.append([x.cari, x.belge, x.tarih, x.vade, float(x.tutar), float(x.kalan), max(g, 0), dilim(g)])
        for c in (3, 4):
            d.cell(d.max_row, c).number_format = "DD.MM.YYYY"
        for c in (5, 6):
            d.cell(d.max_row, c).number_format = PARA
    for j, w in enumerate((30, 16, 13, 13, 14, 14, 13, 16), 1):
        d.column_dimensions[get_column_letter(j)].width = w
    d.freeze_panes = "A2"
    d.auto_filter.ref = d.dimensions

    b = wb.create_sheet("Bilgi")
    for s in [["Rapor tarihi", rapor_tarihi.strftime("%d.%m.%Y")],
              ["Yöntem", "Tahsilatlar (alacak) aynı carinin en eski açık faturasından başlayarak kapatılır (FIFO)"],
              ["Gecikme", "Rapor tarihi − vade tarihi; vade yoksa belge tarihi + varsayılan vade günü"],
              ["Ağırlıklı ort. gecikme", "Σ(açık tutar × gecikme gün) ÷ vadesi geçmiş toplam"]] + [["Uyarı", u] for u in uyarilar]:
        b.append(s)
    b.column_dimensions["A"].width = 22
    b.column_dimensions["B"].width = 110
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    return {"musteri": musteri, "kalemler": kalemler, "uyarilar": uyarilar, "genel": genel}


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description="Cari hareketlerden FIFO alacak yaşlandırma.")
    ap.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "cari_hareketler.csv")
    ap.add_argument("--tarih", help="Rapor tarihi (varsayılan bugün)")
    ap.add_argument("--vade", type=int, default=0, help="Vade tarihi olmayan faturalar için varsayılan vade (gün)")
    ap.add_argument("--limitler", type=Path, help="Kredi limitleri (.xlsx/.csv: Cari, Kredi Limiti)")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "alacak_yaslandirma.xlsx")
    a = ap.parse_args(argv)
    rt = tarih(a.tarih) if a.tarih else date.today()
    if a.girdi == BURASI / "ornek_veri" / "cari_hareketler.csv":
        rt = rt if a.tarih else date(2026, 9, 30)
        a.limitler = a.limitler or BURASI / "ornek_veri" / "kredi_limitleri.csv"
    s = calistir(a.girdi, rt, a.cikti, a.vade, a.limitler)
    tl = lambda x: f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")  # noqa: E731
    print(f"[OK] {len(s['musteri'])} müşteri, {len(s['kalemler'])} açık fatura · toplam açık {tl(s['genel']['toplam'])} TL · "
          f"vadesi geçmiş {tl(s['genel']['vadesi_gecmis'])} TL")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
