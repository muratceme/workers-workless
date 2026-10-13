import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
B = date(2026, 10, 13)
g = lambda ay, gun, yil=2026: date(yil, ay, gun)   # noqa: E731


def p(ad, giris, cikis=None, kimlik=""):
    return {"kimlik": kimlik, "ad": ad, "anahtar": kimlik or main.katla(ad), "giris": giris, "cikis": cikis}


def b(ad, tur, verilis, bildirilen=None, kimlik=""):
    return {"kimlik": kimlik, "ad": ad, "anahtar": kimlik or main.katla(ad), "tur": tur, "verilis": verilis, "bildirilen": bildirilen, "kullanildi": False}


def tek(per, *bil, **kw):
    r = main.kontrol([per], list(bil), kw.pop("bugun", B), **kw)
    x = r["personel"][0]
    return (x["g"]["durum"], x["g"]["gecikme"], x["c"]["durum"] if x["c"] else None, x["c"]["gecikme"] if x["c"] else None, sorted(u["tur"] for u in r["uyarilar"]))


class KontrolTesti(unittest.TestCase):
    def test_giris(self):
        A = p("A", g(10, 5))
        self.assertEqual(tek(A, b("A", "giris", g(10, 4), g(10, 5))), ("Süresinde", 0, None, None, []))
        self.assertEqual(tek(A, b("a", "giris", g(10, 5), g(10, 5))), ("Geç", 1, None, None, ["İşe giriş bildirgesi geç verilmiş"]))
        self.assertEqual(tek(A, b("A", "giris", g(10, 5), g(10, 5)), giris_once=0), ("Süresinde", 0, None, None, []))
        self.assertEqual(tek(A), ("Yok", 9, None, None, ["İşe giriş bildirgesi yok"]))
        self.assertEqual((tek(p("A", g(10, 20))), tek(p("A", g(10, 14))), tek(p("A", g(10, 13)))),
                         (("Bekliyor", 0, None, None, ["İşe giriş bildirgesi bekliyor"]), ("Bekliyor", 0, None, None, ["İşe giriş bildirgesi bekliyor"]), ("Yok", 1, None, None, ["İşe giriş bildirgesi yok"])))
        self.assertEqual(tek(A, b("A", "giris", g(10, 4), g(10, 6))), ("Süresinde", 0, None, None, ["Tarih uyuşmuyor"]))
        self.assertEqual(tek(A, b("A", "giris", g(10, 4))), ("Süresinde", 0, None, None, []))                       # bildirilen tarih boşsa yalnız veriliş tarihine bakılır

    def test_cikis(self):
        gb = b("A", "giris", g(12, 31, 2024), g(1, 1, 2025))
        A = lambda c: p("A", g(1, 1, 2025), c)   # noqa: E731
        self.assertEqual(tek(A(g(9, 30)), dict(gb), b("A", "cikis", g(10, 10), g(9, 30))), ("Süresinde", 0, "Süresinde", 0, []))
        self.assertEqual(tek(A(g(9, 30)), dict(gb), b("A", "cikis", g(10, 11), g(9, 30))), ("Süresinde", 0, "Geç", 1, ["İşten çıkış bildirgesi geç verilmiş"]))
        self.assertEqual(tek(A(g(9, 30)), dict(gb), b("A", "cikis", g(10, 11), g(9, 30)), cikis_gun=15), ("Süresinde", 0, "Süresinde", 0, []))
        self.assertEqual(tek(A(g(10, 2)), dict(gb)), ("Süresinde", 0, "Yok", 1, ["İşten çıkış bildirgesi yok"]))
        self.assertEqual(tek(A(g(10, 3)), dict(gb)), ("Süresinde", 0, "Bekliyor", 0, ["İşten çıkış bildirgesi bekliyor"]))
        self.assertEqual(main.kontrol([A(g(10, 3))], [dict(gb)], B)["uyarilar"][0]["onem"], "Orta")
        self.assertEqual(main.kontrol([A(g(10, 8))], [dict(gb)], B)["uyarilar"][0]["onem"], "Bilgi")
        self.assertEqual(tek(A(g(9, 30)), dict(gb), b("A", "cikis", g(10, 5), g(9, 29))), ("Süresinde", 0, "Süresinde", 0, ["Tarih uyuşmuyor"]))

    def test_esleme_ve_ozet(self):
        per = [p("Ali", g(1, 1, 2024), g(6, 30, 2024), "111"), p("Ali", g(10, 5), None, "111"), p("Veli", g(10, 1))]
        bil = [b("ALİ", "giris", g(10, 4), g(10, 5), "111"), b("Ali", "giris", g(12, 31, 2023), g(1, 1, 2024), "111"), b("Ali", "cikis", g(7, 5, 2024), g(6, 30, 2024), "111"), b("Ayşe", "giris", g(10, 1), g(10, 2)),
               b("Veli", "giris", g(3, 1), g(3, 2))]
        r = main.kontrol(per, bil, B)
        self.assertEqual([(x["ad"], x["giris"], x["g"]["durum"], x["c"]["durum"] if x["c"] else None, x["sorun"]) for x in r["personel"]],
                         [("Ali", g(1, 1, 2024), "Süresinde", "Süresinde", False), ("Veli", g(10, 1), "Yok", None, True), ("Ali", g(10, 5), "Süresinde", None, False)])
        self.assertEqual((r["say"], r["sorunlu"], sorted((u["tur"], u["kim"]) for u in r["uyarilar"])),
                         ({"giris": {"Süresinde": 2, "Yok": 1}, "cikis": {"Süresinde": 1}}, 1, [("Personel listesinde karşılığı olmayan bildirge", "Ayşe"), ("Personel listesinde karşılığı olmayan bildirge", "Veli"), ("İşe giriş bildirgesi yok", "Veli")]))
        self.assertEqual((main.maskele("12345678950"), main.maskele("S-0042"), main.kimlik_metni(12345678950.0)), ("123******50", "S-0042", "12345678950"))
        with self.assertRaises(ValueError):
            main.kontrol([], [], B)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "s.xlsx"
            r = main.calistir(ORNEK / "personel.csv", cikti, B, ORNEK / "bildirgeler.csv")
            self.assertEqual([(x["ad"][-3:].strip(), x["g"]["durum"], x["c"]["durum"] if x["c"] else None) for x in r["personel"]],
                             [("ltı", "Süresinde", "Yok"), ("Beş", "Süresinde", "Geç"), ("edi", "Süresinde", "Bekliyor"), ("Üç", "Yok", None), ("Bir", "Süresinde", None), ("İki", "Geç", None), ("ört", "Bekliyor", None)])
            self.assertEqual((r["sorunlu"], [(u["onem"], u["tur"]) for u in r["uyarilar"]]),
                             (5, [("Yüksek", "İşe giriş bildirgesi geç verilmiş"), ("Yüksek", "İşe giriş bildirgesi yok"), ("Yüksek", "İşten çıkış bildirgesi geç verilmiş"), ("Yüksek", "İşten çıkış bildirgesi yok"),
                                  ("Orta", "Personel listesinde karşılığı olmayan bildirge"), ("Orta", "Tarih uyuşmuyor"), ("Bilgi", "İşe giriş bildirgesi bekliyor"), ("Bilgi", "İşten çıkış bildirgesi bekliyor")]))
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Bildirge Kontrolü", "Uyarılar"])
            self.assertNotIn("12345678950", str([[c.value for c in row] for ws in wb for row in ws.iter_rows()]))

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            py, by = Path(tmp) / "p.csv", Path(tmp) / "b.csv"
            py.write_text("Sicil No;Personel;Giriş;Çıkış\nS1;Ali;05.10.2026;\nS2;Veli;x;\nS3;Can;05.10.2026;01.10.2026\n;Ece;01.10.2026;x\n", encoding="utf-8")
            by.write_text("Sicil;Ad Soyad;Belge Türü;İşlem Tarihi;Tarih\nS1;Ali;İşe Giriş Bildirgesi;04.10.2026;05.10.2026\nS1;Ali;Diğer;04.10.2026;\nS1;Ali;Çıkış;x;\n", encoding="utf-8")
            per, bil, uy = main.oku(py, by)
            self.assertEqual(([(x["anahtar"], x["giris"]) for x in per], [(x["anahtar"], x["tur"], x["verilis"], x["bildirilen"]) for x in bil], sorted(u["tur"] for u in uy)),
                             ([("S1", g(10, 5))], [("S1", "giris", g(10, 4), g(10, 5))], ["Bildirge satırı okunamadı"] * 2 + ["Personel satırı okunamadı"] * 3))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--personel", str(py), "--cikti", c]
            self.assertEqual(main.main(arg + ["--bildirgeler", str(by), "--tarih", "13.10.2026", "--giris-once", "0", "--cikis-gun", "10"]), 0)
            self.assertEqual(main.main(arg), 0)
            for ek in (["--tarih", "x"], ["--giris-once", "-1"]):
                self.assertEqual(main.main(arg + ek), 2)
            self.assertEqual(main.main(arg + ["--bildirgeler", str(Path(tmp) / "yok.csv")]), 1)
            py.write_text("Personel;Çıkış\nAli;1\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
