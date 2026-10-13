"""
SGK Giriş-Çıkış Bildirge Kontrolü — Workers / Workless kod bloğu
İnsan Kaynakları › Bordro ve Özlük İşleri Uzmanı

İşe giriş ve işten çıkış listesini SGK bildirge kayıtlarınızla karşılaştırır; süresinde verilmemiş, eksik veya tarihi uyuşmayan bildirgeleri listeler.
  İşe giriş bildirgesi son günü  = işe giriş tarihi − --giris-once gün (varsayılan 1: en geç işe başlamadan bir gün önce;  0: en geç işe başlanan gün)
  İşten ayrılış bildirgesi son günü = işten çıkış tarihi + --cikis-gun gün (varsayılan 10)
  Durum                          = Süresinde · Geç (bildirge tarihi son günden sonra) · Yok (son gün geçti, kayıt yok) · Bekliyor (son gün gelmedi, kayıt yok)
  Tarih uyuşmazlığı              = bildirgede yazan işe giriş / çıkış tarihi personel kaydındakinden farklı
  Eşleştirme                     = kimlik veya sicil numarasıyla (yoksa ad soyadla); aynı kişinin birden çok çalışma dönemi tarihle eşlenir
Kod SGK sistemine bağlanmaz; yalnız verdiğiniz iki listeyi karşılaştırır. Son gün hafta sonu veya resmi tatile rastlıyorsa süre uzayabilir; kod uzatmaz (ihtiyatlı tarafta kalır).
Giriş süresinin istisnaları (inşaat, balıkçılık, tarım işyerleri, yabancı ülkeye gönderilenler, ilk kez tescil edilen işyeri vb.) için --giris-once değerini siz belirlersiniz.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


PERSONEL = {"kimlik": ("tc kimlik no", "tckn", "tc", "kimlik no", "sicil no", "sicil"), "ad": ("ad soyad", "adi soyadi", "personel", "calisan"), "giris": ("ise giris tarihi", "ise giris", "giris tarihi", "giris"),
            "cikis": ("isten cikis tarihi", "isten cikis", "cikis tarihi", "cikis")}
BILDIRGE = {"kimlik": PERSONEL["kimlik"], "ad": PERSONEL["ad"], "tur": ("tur", "bildirge turu", "belge turu"), "verilis": ("bildirge tarihi", "verilis tarihi", "tescil tarihi", "onay tarihi", "islem tarihi"),
            "bildirilen": ("bildirilen tarih", "ise giris cikis tarihi", "olay tarihi", "tarih")}
TUR = {"giris": "giris", "ise giris": "giris", "ise giris bildirgesi": "giris", "cikis": "cikis", "isten cikis": "cikis", "isten ayrilis": "cikis", "isten ayrilis bildirgesi": "cikis", "ayrilis": "cikis"}
RENK_D = {"Süresinde": YESIL, "Geç": KIRMIZI, "Yok": KIRMIZI, "Bekliyor": SARI}
GUN = timedelta(days=1)


def kimlik_metni(x) -> str:
    if isinstance(x, float) and x == int(x):
        x = int(x)
    return re.sub(r"\s", "", metin(x))


def maskele(k: str) -> str:
    return k[:3] + "*" * (len(k) - 5) + k[-2:] if re.fullmatch(r"\d{11}", k) else k


def oku(personel_yol: Path, bildirge_yol: Path | None = None) -> tuple[list[dict], list[dict], list[dict]]:
    uy, personel, bildirgeler = [], [], []
    for r in kayitlar(personel_yol, PERSONEL, ("ad", "giris")):
        ad, k = metin(r["ad"]), kimlik_metni(r["kimlik"])
        if not ad and not k:
            continue
        g, c = tarih(r["giris"]), tarih(r["cikis"])
        if g is None or (c is None and metin(r["cikis"])) or (c and c < g):
            uy.append(uyari("Yüksek", "Personel satırı okunamadı", f"{ad} (satır {r['_satir']})", f"İşe giriş '{metin(r['giris'])}', işten çıkış '{metin(r['cikis'])}'. Çıkış girişten önce olamaz. Satır kontrol edilmedi."))
            continue
        personel.append({"kimlik": k, "ad": ad, "anahtar": k or katla(ad), "giris": g, "cikis": c})
    if bildirge_yol:
        for r in kayitlar(bildirge_yol, BILDIRGE, ("tur", "verilis")):
            ad, k = metin(r["ad"]), kimlik_metni(r["kimlik"])
            if not ad and not k:
                continue
            t, v, b = TUR.get(katla(r["tur"])), tarih(r["verilis"]), tarih(r["bildirilen"])
            if t is None or v is None or (b is None and metin(r["bildirilen"])):
                uy.append(uyari("Yüksek", "Bildirge satırı okunamadı", f"{ad or maskele(k)} (satır {r['_satir']})", f"Tür '{metin(r['tur'])}', bildirge tarihi '{metin(r['verilis'])}', bildirilen tarih '{metin(r['bildirilen'])}'. Satır alınmadı; bildirge eksik görünebilir."))
                continue
            bildirgeler.append({"kimlik": k, "ad": ad, "anahtar": k or katla(ad), "tur": t, "verilis": v, "bildirilen": b, "kullanildi": False})
    return personel, bildirgeler, uy


def esle(bildirgeler: list[dict], anahtar: str, tur: str, olay: date) -> dict | None:
    aday = [b for b in bildirgeler if b["anahtar"] == anahtar and b["tur"] == tur and not b["kullanildi"]]
    sec = next((b for b in aday if b["bildirilen"] == olay), None) or (min(aday, key=lambda b: abs(((b["bildirilen"] or b["verilis"]) - olay).days)) if aday else None)
    if sec and sec["bildirilen"] != olay and abs(((sec["bildirilen"] or sec["verilis"]) - olay).days) > 31:
        return None                                                                              # başka bir çalışma dönemine ait olabilir
    if sec:
        sec["kullanildi"] = True
    return sec


def kontrol(personel: list[dict], bildirgeler: list[dict], bugun: date, giris_once: int = 1, cikis_gun: int = 10) -> dict:
    if not personel:
        raise ValueError("Kontrol edilecek personel yok.")
    for b in bildirgeler:
        b["kullanildi"] = False
    uy, liste = [], []
    for p in sorted(personel, key=lambda p: (p["giris"], p["ad"])):
        kim = p["ad"] or maskele(p["kimlik"])

        def degerlendir(tur: str, ad: str, olay: date, son: date) -> dict:
            b = esle(bildirgeler, p["anahtar"], tur, olay)
            if b is None:
                durum = "Yok" if bugun > son else "Bekliyor"
                if durum == "Yok":
                    uy.append(uyari("Yüksek", f"{ad} bildirgesi yok", kim, f"{ad} tarihi {olay:%d.%m.%Y}; bildirgenin son günü {son:%d.%m.%Y} idi ({(bugun - son).days} gün geçti). Bildirgeyi hemen verin; idari para cezası riski var."))
                else:
                    kalan = (son - bugun).days
                    uy.append(uyari("Orta" if kalan <= 3 else "Bilgi", f"{ad} bildirgesi bekliyor", kim, f"{ad} tarihi {olay:%d.%m.%Y}; bildirgenin son günü {son:%d.%m.%Y} ({kalan} gün kaldı)."))
                return {"durum": durum, "son": son, "verilis": None, "gecikme": max((bugun - son).days, 0)}
            gec = (b["verilis"] - son).days
            if gec > 0:
                uy.append(uyari("Yüksek", f"{ad} bildirgesi geç verilmiş", kim, f"{ad} tarihi {olay:%d.%m.%Y}, son gün {son:%d.%m.%Y}; bildirge {b['verilis']:%d.%m.%Y} tarihinde verilmiş ({gec} gün geç)."))
            if b["bildirilen"] and b["bildirilen"] != olay:
                uy.append(uyari("Orta", "Tarih uyuşmuyor", kim, f"Personel kaydında {ad.lower()} {olay:%d.%m.%Y}, bildirgede {b['bildirilen']:%d.%m.%Y}. Hangisi doğruysa diğerini düzeltin."))
            return {"durum": "Geç" if gec > 0 else "Süresinde", "son": son, "verilis": b["verilis"], "gecikme": max(gec, 0), "uyusmuyor": bool(b["bildirilen"] and b["bildirilen"] != olay)}

        g = degerlendir("giris", "İşe giriş", p["giris"], p["giris"] - giris_once * GUN)
        c = degerlendir("cikis", "İşten çıkış", p["cikis"], p["cikis"] + cikis_gun * GUN) if p["cikis"] else None
        liste.append({"ad": p["ad"], "kimlik": maskele(p["kimlik"]), "giris": p["giris"], "cikis": p["cikis"], "g": g, "c": c, "sorun": g["durum"] in ("Geç", "Yok") or g.get("uyusmuyor", False) or bool(c and (c["durum"] in ("Geç", "Yok") or c.get("uyusmuyor", False)))})
    for b in bildirgeler:
        if not b["kullanildi"]:
            uy.append(uyari("Orta", "Personel listesinde karşılığı olmayan bildirge", b["ad"] or maskele(b["kimlik"]),
                            f"{'İşe giriş' if b['tur'] == 'giris' else 'İşten çıkış'} bildirgesi {b['verilis']:%d.%m.%Y} tarihinde verilmiş ama personel listesinde eşleşen kayıt yok. İşe başlamayan kişi için verildiyse iptalini değerlendirin."))
    say = {"giris": defaultdict(int), "cikis": defaultdict(int)}
    for x in liste:
        say["giris"][x["g"]["durum"]] += 1
        if x["c"]:
            say["cikis"][x["c"]["durum"]] += 1
    return {"personel": liste, "say": {k: dict(v) for k, v in say.items()}, "sorunlu": sum(1 for x in liste if x["sorun"]), "bugun": bugun, "ayar": (giris_once, cikis_gun), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    go, cg = r["ayar"]
    ws = wb.active
    ws.title = "Bildirge Kontrolü"
    _baslik(ws, ["Ad Soyad", "Kimlik / Sicil", "İşe Giriş", "Giriş Bildirgesi Son Günü", "Giriş Bildirge Tarihi", "Giriş Durumu", "Gecikme (gün)", "İşten Çıkış", "Çıkış Bildirgesi Son Günü", "Çıkış Bildirge Tarihi", "Çıkış Durumu", "Gecikme (gün)",
                 "Not (siz doldurun)"], (28, 16, 12, 14, 14, 12, 9, 12, 14, 14, 12, 9, 28))
    for x in r["personel"]:
        g, c = x["g"], x["c"] or {"son": None, "verilis": None, "durum": "", "gecikme": 0}
        _satir(ws, [x["ad"], x["kimlik"], x["giris"], g["son"], g["verilis"], g["durum"], g["gecikme"] or None, x["cikis"], c["son"], c["verilis"], c["durum"], c["gecikme"] or None, ""], tarih_s=(3, 4, 5, 8, 9, 10),
               dolgu={6: RENK_D.get(g["durum"]), 11: RENK_D.get(c["durum"])})
    ws.append([])
    ws.append([f"Rapor tarihi {r['bugun']:%d.%m.%Y}. Giriş bildirgesi en geç işe girişten {go} gün önce, çıkış bildirgesi en geç çıkıştan sonraki {cg}. gün verilmiş olmalı. Tatil nedeniyle süre uzaması ve sektör istisnaları hesaba katılmaz; kimlik numaraları maskelidir."])
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(personel_yol: Path, cikti: Path, bugun: date, bildirge_yol: Path | None = None, **ayar) -> dict:
    personel, bildirgeler, uy = oku(personel_yol, bildirge_yol)
    r = kontrol(personel, bildirgeler, bugun, **ayar)
    uy += r["uyarilar"]
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="İşe giriş ve çıkış listesini SGK bildirge kayıtlarıyla karşılaştırır.")
    p.add_argument("--personel", type=Path, default=ORNEK / "personel.csv", help="TC Kimlik No (veya Sicil No), Ad Soyad, İşe Giriş Tarihi, İşten Çıkış Tarihi")
    p.add_argument("--bildirgeler", type=Path, help="TC Kimlik No, Ad Soyad, Tür (Giriş / Çıkış), Bildirge Tarihi, Bildirilen Tarih")
    p.add_argument("--tarih", help="Rapor tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--giris-once", type=int, default=1, help="Giriş bildirgesi işe girişten en geç kaç gün önce verilmeli (varsayılan 1; aynı gün için 0)")
    p.add_argument("--cikis-gun", type=int, default=10, help="Çıkış bildirgesi çıkıştan sonra en geç kaç gün içinde verilmeli (varsayılan 10)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "sgk_bildirge_kontrol.xlsx")
    a = p.parse_args(argv)
    ornek = a.personel == ORNEK / "personel.csv"
    bil = a.bildirgeler or (ORNEK / "bildirgeler.csv" if ornek else None)
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 13) if ornek else date.today()
    if bugun is None or a.giris_once < 0 or a.cikis_gun < 0:
        print("[X] --tarih GG.AA.YYYY biçiminde; --giris-once ve --cikis-gun sıfır veya pozitif olmalı")
        return 2
    for y in (a.personel, bil):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.personel, a.cikti, bugun, bil, giris_once=a.giris_once, cikis_gun=a.cikis_gun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(r['personel'])} personel · {r['sorunlu']} kayıtta sorun · giriş: {r['say']['giris']} · çıkış: {r['say']['cikis']}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
