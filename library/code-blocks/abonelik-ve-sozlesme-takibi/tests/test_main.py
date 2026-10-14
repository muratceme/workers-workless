import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 14)


def s(ad, bitis_gun=None, ihbar=0, oto=False, aylik=100, kategori="K", pb="TRY", bas=None, artis=""):
    return {"ad": ad, "kategori": kategori, "tedarikci": "", "abone": "", "bas": bas, "bitis": BUGUN + timedelta(days=bitis_gun) if bitis_gun is not None else None, "tutar": D(aylik),
            "aylik": D(aylik), "pb": pb, "oto": oto, "ihbar": ihbar, "artis": artis, "sorumlu": "", "lokasyon": ""}


class AnalizTesti(unittest.TestCase):
    def test_durum(self):
        l = [s("a", -1), s("b", -1, oto=True), s("c", 40, 30, True), s("d", 20, 30, True), s("e", 50), s("f", 200, 30, True), s("g"), s("h", 100, 30, False, bas=date(2020, 10, 20), artis="TÜFE")]
        r = main.analiz(l, BUGUN)
        d = {x["ad"]: x["durum"] for x in r["sozlesmeler"]}
        self.assertEqual(d, {"a": "Süresi doldu", "b": "Süresi doldu", "c": "İhbar yaklaşıyor", "d": "İhbar süresi geçti", "e": "Bitiş yaklaşıyor", "f": "Devam ediyor", "g": "Süresiz",
                             "h": "Devam ediyor"})
        x = {v["ad"]: v for v in r["sozlesmeler"]}
        self.assertEqual((x["c"]["ihbar_son"], x["c"]["ihbar_kalan"], x["e"]["ihbar_son"], x["h"]["artis_tarihi"]), (date(2026, 10, 24), 10, None, date(2026, 10, 20)))
        self.assertEqual(r["kategoriler"], [{"kategori": "K", "pb": "TRY", "adet": 7, "aylik": D(700), "yillik": D(8400)}])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Bilgi", "Fiyat artış dönemi", "h"), ("Bilgi", "İhbar süresi geçti", "d"), ("Orta", "Bitiş yaklaşıyor", "e"), ("Orta", "Fesih ihbarı son günü yaklaşıyor", "c"),
                          ("Orta", "Süresi doldu", "b"), ("Yüksek", "Süresi doldu", "a")])
        with self.assertRaises(ValueError):
            main.analiz([], BUGUN)

    def test_yildonumu(self):
        self.assertEqual([main.yildonumu(date(2020, 10, 14), BUGUN), main.yildonumu(date(2020, 1, 1), BUGUN), main.yildonumu(date(2024, 2, 29), date(2026, 10, 14))],
                         [date(2026, 10, 14), date(2027, 1, 1), date(2027, 3, 1)])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "sozlesmeler.csv", cikti, BUGUN)
            o = r["ozet"]
            self.assertEqual((len(r["sozlesmeler"]), o["Süresi doldu"], o["İhbar yaklaşıyor"], o["Bitiş yaklaşıyor"], o["İhbar süresi geçti"], o["Süresiz"]), (13, 1, 2, 2, 1, 2))
            asansor = next(x for x in r["sozlesmeler"] if x["ad"] == "Asansör bakım")
            self.assertEqual(asansor["aylik"], D(4500))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Sözleşmeler", "Maliyet", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "s.csv"
            y.write_text("Hizmet;Bedel;Periyot;Başlangıç;Bitiş;İhbar Süresi\nA;1.200;Yıllık;;;\nB;100;Haftalık;;;\nC;100;;01.01.2026;01.01.2025;\nD;x;;;;\nE;300;3 aylık;;;x\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["ad"], x["aylik"]) for x in l], [u["tur"] for u in uy]), ([("A", D(100))], ["Satır okunamadı", "Bitiş başlangıçtan önce", "Satır okunamadı", "Satır okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--sozlesmeler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "14.10.2026", "--ihbar", "45", "--bitis", "90"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + x) for x in (["--tarih", "x"], ["--ihbar", "-1"])], [2, 2])
            self.assertEqual(main.main(["--sozlesmeler", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Hizmet;Bedel\nA;x\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
