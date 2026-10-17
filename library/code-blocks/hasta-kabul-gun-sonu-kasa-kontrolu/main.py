"""
Hasta Kabul Gün Sonu Kasa Kontrolü — Gerisi Bende kod bloğu
Hasta Hizmetleri › Hasta Kayıt Kabul Görevlisi
Gün içindeki tahsilat, fatura ve iadeleri HBYS kayıtlarıyla karşılaştırıp kasa devir raporunu hazırlar.
  Beklenen nakit   = önceki günden devir + nakit tahsilat − nakit iade − bankaya / merkez kasaya teslim edilen
  Nakit farkı      = sayılan nakit − beklenen nakit
  Kart farkı       = POS gün sonu toplamı − (kartla tahsilat − karta iade)
  Sonraki güne devir = sayılan nakit (sayım girilmemişse beklenen nakit)
  Fatura kontrolü  = faturadaki hasta payı (tahsil edilecek) − o faturaya yapılan net tahsilat (tahsilat − iade)
  Kontroller       = mükerrer makbuz · faturasız tahsilat · hastaya yapılan iade o hastadan yapılan tahsilattan fazla · iptal edilen makbuzlar
İptal edilen makbuzlar kasa toplamlarına girmez; ayrıca listelenir.
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

HAREKET = {"makbuz": ("makbuz no", "makbuz", "islem no", "belge no", "dekont no"), "saat": ("saat", "islem saati", "tarih saat", "tarih"), "protokol": ("protokol no", "protokol", "hasta no", "dosya no"),
           "hasta": ("hasta", "hasta adi", "ad soyad"), "islem": ("islem", "islem turu", "hareket turu", "tur"), "odeme": ("odeme turu", "odeme sekli", "odeme", "tahsilat turu"),
           "tutar": ("tutar", "islem tutari", "toplam"), "fatura": ("fatura no", "fatura", "fatura numarasi"), "kullanici": ("kullanici", "veznedar", "personel", "islem yapan")}
FATURA = {"fatura": ("fatura no", "fatura", "fatura numarasi"), "protokol": ("protokol no", "protokol", "hasta no", "dosya no"), "hasta": ("hasta", "hasta adi", "ad soyad"),
          "tutar": ("tahsil edilecek", "hasta payi", "hasta odemesi", "odenecek tutar", "fatura tutari", "tutar")}
SIFIR = Decimal(0)
TURLER = ("Nakit", "Kart", "Havale / EFT")


def _islem(x) -> str | None:
    k = katla(x)
    if not k or "tahsil" in k or k in ("giris", "odeme", "satis"):
        return "Tahsilat"
    if "iade" in k:
        return "İade"
    if "iptal" in k:
        return "İptal"
    return None


def _odeme(x) -> str:
    k = katla(x)
    if not k or "nakit" in k:
        return "Nakit"
    if any(s in k for s in ("kart", "kredi", "pos", "visa", "master")):
        return "Kart"
    if any(s in k for s in ("havale", "eft", "fast", "banka")):
        return "Havale / EFT"
    return metin(x)


def oku_hareket(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, HAREKET, ("makbuz", "tutar")):
        makbuz = metin(r["makbuz"])
        if not makbuz:
            continue
        tutar, islem = para(r["tutar"]), _islem(r["islem"])
        if tutar is None or islem is None:
            uy.append(uyari("Yüksek", "Hareket satırı okunamadı", makbuz, f"Satır {r['_satir']}: tutar ya da işlem türü ('{metin(r['islem'])}') okunamadı. Satır alınmadı."))
            continue
        if katla(makbuz) in gorulen:
            uy.append(uyari("Yüksek", "Makbuz no iki kez", makbuz, f"Satır {r['_satir']}: aynı makbuz numarası ikinci kez geçiyor; mükerrer kayıt olabilir. İkinci satır alınmadı."))
            continue
        gorulen.add(katla(makbuz))
        l.append({"makbuz": makbuz, "saat": metin(r["saat"]), "protokol": metin(r["protokol"]), "hasta": metin(r["hasta"]), "islem": islem, "odeme": _odeme(r["odeme"]), "tutar": abs(tutar),
                  "fatura": metin(r["fatura"]), "kullanici": metin(r["kullanici"]) or "-"})
    return l, uy


def oku_fatura(yol: Path) -> tuple[dict, list[dict]]:
    uy, d = [], {}
    for r in kayitlar(yol, FATURA, ("fatura", "tutar")):
        no = metin(r["fatura"])
        if not no:
            continue
        tutar = para(r["tutar"])
        if tutar is None or tutar < 0:
            uy.append(uyari("Yüksek", "Fatura satırı okunamadı", no, f"Satır {r['_satir']}: tahsil edilecek tutar okunamadı. Fatura alınmadı."))
            continue
        if katla(no) in d:
            uy.append(uyari("Orta", "Fatura iki kez yazılmış", no, f"Satır {r['_satir']}: ikinci satır alınmadı."))
            continue
        d[katla(no)] = {"fatura": no, "protokol": metin(r["protokol"]), "hasta": metin(r["hasta"]), "tutar": tutar}
    return d, uy


def kontrol(hareketler: list[dict], faturalar: dict | None = None, devir: Decimal = SIFIR, bankaya: Decimal = SIFIR, sayim: Decimal | None = None, pos: Decimal | None = None,
            tolerans: Decimal = Decimal(1)) -> dict:
    if not hareketler:
        raise ValueError("Kasa hareketi yok.")
    uy = []
    gecerli = [h for h in hareketler if h["islem"] != "İptal"]
    top = lambda islem, odeme=None: sum((h["tutar"] for h in gecerli if h["islem"] == islem and (odeme is None or h["odeme"] == odeme)), SIFIR)   # noqa: E731
    turler = list(TURLER) + sorted({h["odeme"] for h in gecerli} - set(TURLER))
    odeme = [{"tur": t, "tahsilat": top("Tahsilat", t), "iade": top("İade", t), "net": top("Tahsilat", t) - top("İade", t), "adet": sum(1 for h in gecerli if h["odeme"] == t)} for t in turler]
    nakit, kart = odeme[0], odeme[1]
    beklenen = devir + nakit["net"] - bankaya
    fark = sayim - beklenen if sayim is not None else None
    pos_fark = pos - kart["net"] if pos is not None else None
    if fark is not None and abs(fark) > tolerans:
        uy.append(uyari("Yüksek", "Kasa eksik" if fark < 0 else "Kasa fazla", "Nakit", f"Sayılan {sayim:,.2f}, beklenen {beklenen:,.2f}; fark {fark:,.2f} (tolerans {tolerans:,.2f})."))
    if pos_fark is not None and abs(pos_fark) > tolerans:
        uy.append(uyari("Yüksek", "POS farkı", "Kart", f"POS gün sonu {pos:,.2f}, HBYS kart neti {kart['net']:,.2f}; fark {pos_fark:,.2f}."))
    if sayim is None:
        uy.append(uyari("Bilgi", "Sayım girilmedi", "Nakit", "Nakit farkı hesaplanmadı; devir beklenen nakit üzerinden yazıldı."))
    for h in hareketler:
        if h["islem"] == "İptal":
            uy.append(uyari("Bilgi", "İptal edilen makbuz", h["makbuz"], f"{h['tutar']:,.2f} {h['odeme']} · protokol {h['protokol'] or '-'} · {h['kullanici']}. İptal gerekçesi ve onayı kontrol edilmeli."))
        elif h["islem"] == "Tahsilat" and not h["fatura"]:
            uy.append(uyari("Orta", "Faturasız tahsilat", h["makbuz"], f"{h['tutar']:,.2f} {h['odeme']} · protokol {h['protokol'] or '-'}: fatura numarası yok."))
    hasta = defaultdict(lambda: [SIFIR, SIFIR])
    for h in gecerli:
        if h["protokol"]:
            hasta[katla(h["protokol"])][0 if h["islem"] == "Tahsilat" else 1] += h["tutar"]
    adlar = {katla(h["protokol"]): h["protokol"] for h in gecerli if h["protokol"]}
    for k, (t, i) in sorted(hasta.items()):
        if i > t:
            uy.append(uyari("Yüksek", "İade tahsilattan fazla", adlar[k], f"Bugün bu protokolden {t:,.2f} tahsilat, {i:,.2f} iade. Önceki günlere ait tahsilatın iadesiyse belgesini kontrol edin."))
    fkontrol = []
    if faturalar is not None:
        net = defaultdict(lambda: SIFIR)
        for h in gecerli:
            if h["fatura"]:
                net[katla(h["fatura"])] += h["tutar"] if h["islem"] == "Tahsilat" else -h["tutar"]
        for k, f in sorted(faturalar.items()):
            fk = {**f, "tahsil": net.get(k, SIFIR), "fark": f["tutar"] - net.get(k, SIFIR)}
            fkontrol.append(fk)
            if abs(fk["fark"]) > tolerans:
                uy.append(uyari("Orta", "Fatura tahsilatı eksik" if fk["fark"] > 0 else "Fatura tahsilatı fazla", f["fatura"],
                                f"Protokol {f['protokol'] or '-'}: tahsil edilecek {f['tutar']:,.2f}, tahsil edilen {fk['tahsil']:,.2f}; fark {fk['fark']:,.2f}."))
        for h in gecerli:
            if h["islem"] == "Tahsilat" and h["fatura"] and katla(h["fatura"]) not in faturalar:
                uy.append(uyari("Orta", "Faturası listede yok", h["makbuz"], f"{h['tutar']:,.2f} tahsilatın fatura numarası ({h['fatura']}) fatura listesinde yok."))
    kullanicilar = []
    for k in sorted({h["kullanici"] for h in hareketler}, key=katla):
        l = [h for h in hareketler if h["kullanici"] == k]
        kullanicilar.append({"kullanici": k, "adet": len(l), "tahsilat": sum((h["tutar"] for h in l if h["islem"] == "Tahsilat"), SIFIR),
                             "iade": sum((h["tutar"] for h in l if h["islem"] == "İade"), SIFIR), "iptal": sum(1 for h in l if h["islem"] == "İptal"),
                             "nakit_net": sum((h["tutar"] if h["islem"] == "Tahsilat" else -h["tutar"] for h in l if h["islem"] != "İptal" and h["odeme"] == "Nakit"), SIFIR)})
    return {"odeme": odeme, "devir": devir, "bankaya": bankaya, "beklenen": beklenen, "sayim": sayim, "fark": fark, "pos": pos, "pos_fark": pos_fark,
            "sonraki": sayim if sayim is not None else beklenen, "kullanicilar": kullanicilar, "faturalar": fkontrol, "hareketler": hareketler,
            "toplam_tahsilat": top("Tahsilat"), "toplam_iade": top("İade"), "iptal": sum(1 for h in hareketler if h["islem"] == "İptal"), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, gun: str, tolerans: Decimal) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kasa Devir Raporu"
    _baslik(ws, ["Kalem", "Tutar"], (44, 16))
    nakit, kart = r["odeme"][0], r["odeme"][1]
    satirlar = [("Önceki günden devir", r["devir"]), ("+ Nakit tahsilat", nakit["tahsilat"]), ("− Nakit iade", nakit["iade"]), ("− Bankaya / merkez kasaya teslim", r["bankaya"]),
                ("= Beklenen nakit", r["beklenen"]), ("Sayılan nakit", r["sayim"]), ("Nakit farkı", r["fark"]), (None, None), ("Kartla tahsilat", kart["tahsilat"]), ("− Karta iade", kart["iade"]),
                ("= HBYS kart neti", kart["net"]), ("POS gün sonu", r["pos"]), ("Kart farkı", r["pos_fark"]), (None, None), ("Toplam tahsilat (tüm ödeme türleri)", r["toplam_tahsilat"]),
                ("Toplam iade", r["toplam_iade"]), ("İptal edilen makbuz sayısı", r["iptal"]), (None, None), ("SONRAKİ GÜNE DEVREDEN NAKİT", r["sonraki"])]
    for ad, v in satirlar:
        if ad is None:
            ws.append([])
            continue
        renk = None
        if ad in ("Nakit farkı", "Kart farkı") and v is not None:
            renk = YESIL if abs(v) <= tolerans else KIRMIZI
        k = _satir(ws, [ad, v], para_s=(2,) if ad != "İptal edilen makbuz sayısı" else (), dolgu={2: renk})
        if ad.startswith("=") or ad.isupper():
            for h in ws[k]:
                h.font = Font(bold=True)
    ws.append([])
    ws.append([f"Gün: {gun}" if gun else "Gün belirtilmedi.", None])
    ws.append(["Teslim eden: ................................", "Teslim alan: ................................"])
    ws = wb.create_sheet("Ödeme Türleri")
    _baslik(ws, ["Ödeme Türü", "İşlem", "Tahsilat", "İade", "Net"], (16, 8, 14, 14, 14))
    for o in r["odeme"]:
        _satir(ws, [o["tur"], o["adet"], o["tahsilat"], o["iade"], o["net"]], para_s=(3, 4, 5))
    ws = wb.create_sheet("Kullanıcılar")
    _baslik(ws, ["Kullanıcı", "İşlem", "Tahsilat", "İade", "İptal", "Nakit Net"], (18, 8, 14, 14, 8, 14))
    for k in r["kullanicilar"]:
        _satir(ws, [k["kullanici"], k["adet"], k["tahsilat"], k["iade"], k["iptal"], k["nakit_net"]], para_s=(3, 4, 6), dolgu={5: SARI if k["iptal"] else None})
    if r["faturalar"]:
        ws = wb.create_sheet("Fatura Kontrolü")
        _baslik(ws, ["Fatura No", "Protokol No", "Hasta", "Tahsil Edilecek", "Tahsil Edilen", "Fark"], (14, 13, 22, 14, 14, 12))
        for f in r["faturalar"]:
            _satir(ws, [f["fatura"], f["protokol"], f["hasta"], f["tutar"], f["tahsil"], f["fark"]], para_s=(4, 5, 6), dolgu={6: KIRMIZI if abs(f["fark"]) > tolerans else None})
    ws = wb.create_sheet("Hareketler")
    _baslik(ws, ["Makbuz No", "Saat", "Protokol No", "Hasta", "İşlem", "Ödeme Türü", "Tutar", "Fatura No", "Kullanıcı"], (12, 10, 12, 22, 9, 12, 12, 12, 14), "B2")
    for h in r["hareketler"]:
        _satir(ws, [h["makbuz"], h["saat"], h["protokol"], h["hasta"], h["islem"], h["odeme"], h["tutar"], h["fatura"], h["kullanici"]], para_s=(7,),
               dolgu={5: MAVI if h["islem"] == "İade" else SARI if h["islem"] == "İptal" else None})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(hareket_yol: Path, cikti: Path, fatura_yol: Path | None = None, devir: Decimal = SIFIR, bankaya: Decimal = SIFIR, sayim: Decimal | None = None, pos: Decimal | None = None,
             tolerans: Decimal = Decimal(1), gun: str = "") -> dict:
    hareketler, uy = oku_hareket(hareket_yol)
    faturalar, uy2 = oku_fatura(fatura_yol) if fatura_yol else (None, [])
    r = kontrol(hareketler, faturalar, devir, bankaya, sayim, pos, tolerans)
    r["uyarilar"] = sirala(uy + uy2 + r["uyarilar"])
    rapor_yaz(cikti, r, gun, tolerans)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Gün içindeki tahsilat, fatura ve iadeleri HBYS kayıtlarıyla karşılaştırıp kasa devir raporunu hazırlar.")
    p.add_argument("--hareketler", type=Path, help="HBYS kasa hareketleri: Makbuz No, Saat, Protokol No, Hasta, İşlem, Ödeme Türü, Tutar, Fatura No, Kullanıcı")
    p.add_argument("--faturalar", type=Path, help="Günün faturaları (isteğe bağlı): Fatura No, Protokol No, Hasta, Tahsil Edilecek")
    p.add_argument("--devir", default="0", help="Önceki günden devreden nakit")
    p.add_argument("--bankaya", default="0", help="Gün içinde bankaya ya da merkez kasaya teslim edilen nakit")
    p.add_argument("--sayim", help="Gün sonunda sayılan nakit")
    p.add_argument("--pos", help="POS cihazlarının gün sonu toplamı")
    p.add_argument("--tolerans", default="1", help="Bu tutara kadar fark uyarı vermez (varsayılan 1)")
    p.add_argument("--gun", default="", help="Raporun günü (yalnız başlıkta görünür), ör. 30.09.2026")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "hasta_kabul_kasa_devir.xlsx")
    a = p.parse_args(argv)
    ornek = a.hareketler is None
    hy = a.hareketler or ORNEK / "kasa_hareketleri.csv"
    fy = a.faturalar or (ORNEK / "gunun_faturalari.csv" if ornek else None)
    if ornek and a.sayim is None and a.pos is None:
        a.devir, a.bankaya, a.sayim, a.pos, a.gun = "2.500", "10.000", "6.665", "19.450", a.gun or "30.09.2026"
    sayilar = [para(x) for x in (a.devir, a.bankaya, a.tolerans)]
    opsiyon = [para(x) if x is not None else None for x in (a.sayim, a.pos)]
    if any(x is None or x < 0 for x in sayilar) or any(x is None for x, ham in zip(opsiyon, (a.sayim, a.pos)) if ham is not None):
        print("[X] --devir, --bankaya, --tolerans sıfır ya da üstü; --sayim ve --pos sayı olmalı")
        return 2
    for y in (hy, fy):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(hy, a.cikti, fy, sayilar[0], sayilar[1], opsiyon[0], opsiyon[1], sayilar[2], a.gun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(r['hareketler'])} hareket · tahsilat {r['toplam_tahsilat']:,.2f} · iade {r['toplam_iade']:,.2f} · iptal {r['iptal']}")
    print(f"     Beklenen nakit {r['beklenen']:,.2f}" + (f" · sayılan {r['sayim']:,.2f} · fark {r['fark']:,.2f}" if r["sayim"] is not None else "")
          + (f" · POS farkı {r['pos_fark']:,.2f}" if r["pos_fark"] is not None else "") + f" · devreden {r['sonraki']:,.2f}")
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
