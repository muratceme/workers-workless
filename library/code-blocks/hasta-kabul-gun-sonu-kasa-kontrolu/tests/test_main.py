import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def h(makbuz, islem, odeme, tutar, fatura="", protokol="", kullanici="A"):
    return {"makbuz": makbuz, "saat": "", "protokol": protokol, "hasta": "", "islem": islem, "odeme": odeme, "tutar": D(tutar), "fatura": fatura, "kullanici": kullanici}


def veri():
    return [h("M1", "Tahsilat", "Nakit", 500, "F1", "P1"), h("M2", "Tahsilat", "Kart", 1200, "F2", "P2"), h("M3", "İade", "Nakit", 100, "F1", "P1", "B"), h("M4", "Tahsilat", "Nakit", 300, "", "P3", "B"),
            h("M5", "İptal", "Nakit", 200, "F9", "P4"), h("M6", "İade", "Kart", 200, "F2", "P2", "B"), h("M7", "İade", "Nakit", 50, "", "P7"), h("M8", "Tahsilat", "Çek", 10, "F5", "P8")]


FATURA = {"f1": {"fatura": "F1", "protokol": "P1", "hasta": "", "tutar": D(400)}, "f2": {"fatura": "F2", "protokol": "P2", "hasta": "", "tutar": D(1200)},
          "f4": {"fatura": "F4", "protokol": "P6", "hasta": "", "tutar": D(250)}}


class KontrolTesti(unittest.TestCase):
    def test_kasa(self):
        r = main.kontrol(veri(), FATURA, D(1000), D(500), D(1145), D(1000))
        self.assertEqual([(o["tur"], o["adet"], o["tahsilat"], o["iade"], o["net"]) for o in r["odeme"]],
                         [("Nakit", 4, D(800), D(150), D(650)), ("Kart", 2, D(1200), D(200), D(1000)), ("Havale / EFT", 0, D(0), D(0), D(0)), ("Çek", 1, D(10), D(0), D(10))])
        self.assertEqual((r["beklenen"], r["fark"], r["pos_fark"], r["sonraki"], r["toplam_tahsilat"], r["toplam_iade"], r["iptal"]), (D(1150), D(-5), D(0), D(1145), D(2010), D(350), 1))
        self.assertEqual([(f["fatura"], f["tahsil"], f["fark"]) for f in r["faturalar"]], [("F1", D(400), D(0)), ("F2", D(1000), D(200)), ("F4", D(0), D(250))])
        self.assertEqual([(k["kullanici"], k["adet"], k["tahsilat"], k["iade"], k["iptal"], k["nakit_net"]) for k in r["kullanicilar"]],
                         [("A", 5, D(1710), D(50), 1, D(450)), ("B", 3, D(300), D(300), 0, D(200))])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Bilgi", "İptal edilen makbuz", "M5"), ("Orta", "Fatura tahsilatı eksik", "F2"), ("Orta", "Fatura tahsilatı eksik", "F4"), ("Orta", "Faturası listede yok", "M8"),
                          ("Orta", "Faturasız tahsilat", "M4"), ("Yüksek", "Kasa eksik", "Nakit"), ("Yüksek", "İade tahsilattan fazla", "P7")])

    def test_secenekler(self):
        r = main.kontrol(veri(), tolerans=D(5))
        self.assertEqual((r["beklenen"], r["fark"], r["pos_fark"], r["sonraki"], r["faturalar"]), (D(650), None, None, D(650), []))
        self.assertEqual(sorted(u["tur"] for u in r["uyarilar"]), ["Faturasız tahsilat", "Sayım girilmedi", "İade tahsilattan fazla", "İptal edilen makbuz"])
        r = main.kontrol(veri(), devir=D(0), sayim=D(700), pos=D(900), tolerans=D(5))
        self.assertEqual(sorted(u["tur"] for u in r["uyarilar"] if u["onem"] == "Yüksek"), ["Kasa fazla", "POS farkı", "İade tahsilattan fazla"])
        self.assertEqual([main._islem(x) for x in ("", "Tahsilat", "İADE", "iptal edildi", "x")], ["Tahsilat", "Tahsilat", "İade", "İptal", None])
        self.assertEqual([main._odeme(x) for x in ("", "NAKİT", "Kredi Kartı", "POS", "EFT", "Çek")], ["Nakit", "Nakit", "Kart", "Kart", "Havale / EFT", "Çek"])
        with self.assertRaises(ValueError):
            main.kontrol([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "kasa_hareketleri.csv", cikti, ORNEK / "gunun_faturalari.csv", D(2500), D(10000), D(6665), D(19450), gun="30.09.2026")
            self.assertEqual((r["beklenen"], r["fark"], r["pos_fark"], r["sonraki"], r["toplam_tahsilat"], r["toplam_iade"]), (D(6670), D(-5), D(0), D(6665), D(39370), D(750)))
            self.assertEqual([(u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Kasa eksik", "Nakit"), ("İade tahsilattan fazla", "P-24402"), ("Fatura tahsilatı eksik", "F-55010"), ("Fatura tahsilatı eksik", "F-55012"),
                              ("Faturasız tahsilat", "M-1011"), ("İptal edilen makbuz", "M-1013")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Kasa Devir Raporu", "Ödeme Türleri", "Kullanıcılar", "Fatura Kontrolü", "Hareketler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, z = Path(tmp) / "h.csv", Path(tmp) / "f.csv"
            y.write_text("İşlem No;Tutar;Hareket Türü;Ödeme Şekli\n1;1.250,50;;\n1;5;;\n2;x;;\n3;5;transfer;\n4;-100;İade;Kart\n;5;;\n", encoding="utf-8")
            z.write_text("Fatura No;Hasta Payı\nF1;100\nf1;5\nF2;x\n", encoding="utf-8")
            l, uy = main.oku_hareket(y)
            self.assertEqual(([(x["makbuz"], x["islem"], x["odeme"], x["tutar"], x["kullanici"]) for x in l], [u["tur"] for u in uy]),
                             ([("1", "Tahsilat", "Nakit", D("1250.50"), "-"), ("4", "İade", "Kart", D(100), "-")], ["Makbuz no iki kez", "Hareket satırı okunamadı", "Hareket satırı okunamadı"]))
            d, uy = main.oku_fatura(z)
            self.assertEqual((list(d), [u["tur"] for u in uy]), (["f1"], ["Fatura iki kez yazılmış", "Fatura satırı okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--hareketler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--faturalar", str(z), "--devir", "1.000", "--bankaya", "0", "--sayim", "2.150,50", "--pos", "-100", "--tolerans", "0", "--gun", "01.10.2026"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--devir", "-1"], ["--sayim", "x"], ["--tolerans", "y"])], [2] * 3)
            self.assertEqual(main.main(arg + ["--faturalar", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("İşlem No;Tutar\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
