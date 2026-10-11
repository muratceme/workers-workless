import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
# baz oran 1.500 / 10.000 = %15
P = {"ozkaynak": D(1500), "rav": D(10000), "yp_rav": D(4000), "yp_poz": D(-200), "nfg": D(-10), "deger": D(-20), "kredi": D(8000), "takip": D(200), "karsilik": D(50), "kar": D(100)}


def senaryo(ad, kur=0, faiz=0, takip=0, karsilik=None):
    return {"ad": ad, "aciklama": "", "kur": D(kur), "faiz": D(faiz), "takip": D(takip), "karsilik": D(karsilik) if karsilik is not None else None}


class FormulTesti(unittest.TestCase):
    def test_birlesik_senaryo(self):
        r = main.uygula(P, D(20), D(200), D(2))
        self.assertEqual((r["kur_etki"], r["nfg_etki"], r["deger_etki"], r["ek_takip"], r["kredi_zarari"]), (D(-40), D(-20), D(-40), D(160), D(80)))
        self.assertEqual((r["sok"], r["kar"], r["sermaye_etki"], r["ozkaynak"], r["rav"]), (D(-180), D(-80), D(-80), D(1420), D(10800)))
        self.assertEqual(main.yuvarla(r["syr"]), D("13.15"))                          # 1.420 / 10.800
        self.assertEqual(main.yuvarla(r["takip_orani"]), D("4.50"))                   # (200 + 160) / 8.000

    def test_kar_tamponu(self):
        self.assertEqual(main.uygula(P, D(0), D(100), D(0))["sermaye_etki"], D(0))    # −30 zarar kârdan karşılanır
        r = main.uygula(P, D(20), D(200), D(2), kar_tamponu=False)
        self.assertEqual((r["sermaye_etki"], main.yuvarla(r["syr"])), (D(-180), D("12.22")))

    def test_olumlu_etki_sermayeye_eklenmez(self):
        r = main.uygula(P, D(-50), D(-500), D(0))                                     # kısa pozisyon + faiz düşüşü kazandırır
        self.assertEqual((r["sok"], r["sermaye_etki"], r["rav"]), (D(250), D(0), D(8000)))
        self.assertEqual(main.yuvarla(r["syr"]), D("18.75"))                          # yalnız RAV küçüldüğü için

    def test_senaryo_karsilik_orani(self):
        self.assertEqual(main.uygula(P, D(0), D(0), D(5), D(80))["kredi_zarari"], D(320))
        self.assertEqual(main.uygula(P, D(0), D(0), D(5))["kredi_zarari"], D(200))

    def test_durum_ve_ihtiyac(self):
        r = main.hesapla(P, [senaryo("Hafif", 5), senaryo("Orta", 40, 500, 3), senaryo("Ağır", 100, 2000, 15, 80)])
        s = {x["ad"]: x for x in r["senaryolar"]}
        self.assertEqual([x["durum"] for x in r["senaryolar"]], ["Yeterli", "Hedefin altında", "Asgarinin altında"])
        # Orta: şok −80 − 50 − 100 − 120 = −350; kâr −250; özkaynak 1.250; RAV 11.600 → %10,78
        self.assertEqual((s["Orta"]["sok"], s["Orta"]["ozkaynak"], s["Orta"]["rav"], main.yuvarla(s["Orta"]["syr"])), (D(-350), D(1250), D(11600), D("10.78")))
        self.assertEqual(main.yuvarla(s["Orta"]["ihtiyac"]), D("142.00"))             # 0,12 × 11.600 − 1.250
        self.assertEqual(s["Orta"]["en_buyuk"], "Kredi zararı")
        self.assertEqual(sorted(u["tur"] for u in r["uyarilar"]), ["Asgarinin altında", "Hedefin altında"])
        self.assertEqual(main.yuvarla(r["baz"]["syr"]), D("15.00"))

    def test_ters_stres(self):
        t = {x["ad"]: x["deger"] for x in main.ters_stres(P, D(8))}
        self.assertEqual(t["Takip oranı artışı"], D("20.0"))                          # (1.600 − 40x) / 10.000 = %8
        self.assertEqual(t["Faiz şoku"], D("2666.7"))                                 # (1.600 − 0,3x) / 10.000 = %8
        self.assertEqual(t["Kur şoku"], D("153.8"))                                   # (1.600 − 2x) / (10.000 + 40x) = %8
        uzun = P | {"yp_poz": D(500), "yp_rav": D(0), "nfg": D(10), "deger": D(0)}
        t = {x["ad"]: x["deger"] for x in main.ters_stres(uzun, D(8))}
        self.assertEqual(t["Kur şoku"], None)                                         # uzun pozisyon: kur en çok %95 düşer, zarar 475 yetmez
        self.assertEqual(t["Faiz şoku"], None)                                        # faiz düşüşü zarar ettirir ama −5.000 bp'de zarar 500: yetmez
        dusuk = main.ters_stres(P | {"ozkaynak": D(700)}, D(8))
        self.assertEqual({x["deger"] for x in dusuk}, {D(0)})


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "s.xlsx"
            r = main.calistir(ORNEK / "parametreler.csv", ORNEK / "senaryolar.csv", cikti)
            s = {x["ad"]: x for x in r["senaryolar"]}
            self.assertEqual(main.yuvarla(r["baz"]["syr"]), D("15.00"))
            self.assertEqual((s["Faiz şoku"]["sok"], s["Faiz şoku"]["sermaye_etki"]), (D(-1630000000), D(-330000000)))     # 10 × (68 + 95) milyon; kâr 1,3 milyar
            self.assertEqual((main.yuvarla(s["Şiddetli (birleşik)"]["syr"]), s["Şiddetli (birleşik)"]["durum"]), (D("9.34"), "Hedefin altında"))
            self.assertEqual(main.yuvarla(s["Kur şoku"]["syr"]), D("13.64"))          # 5,4 / 39,6 milyar
            self.assertEqual({t["ad"]: t["deger"] for t in r["ters"]}, {"Kur şoku": D("262.5"), "Faiz şoku": D("2343.6"), "Takip oranı artışı": D("24.5")})
            self.assertEqual(load_workbook(cikti).sheetnames, ["Senaryo Sonuçları", "Etki Dökümü", "Ters Stres", "Parametreler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            p, s = Path(tmp) / "p.csv", Path(tmp) / "s.csv"
            p.write_text("Alan;Değer\nÖzkaynak;1.500\nRisk Ağırlıklı Varlıklar;10.000\nToplam Krediler;8.000\nKarşılık Oranı (%);50\n", encoding="utf-8")
            s.write_text("Senaryo;Kur Şoku (%);Faiz Şoku (bp);Takip Oranı Artışı (puan)\nA;10;;2\nB;on;100;-1\n", encoding="utf-8")
            prm, uy = main.parametre_oku(p)
            self.assertEqual((prm["karsilik"], prm["yp_poz"], prm["kar"]), (D(50), D(0), D(0)))
            self.assertEqual(len(uy), 6)                                              # girilmeyen 6 parametre
            sn, uy = main.senaryo_oku(s)
            self.assertEqual([(x["kur"], x["faiz"], x["takip"]) for x in sn], [(D(10), D(0), D(2)), (D(0), D(100), D(0))])
            self.assertEqual(sorted(u["tur"] for u in uy), ["Senaryo değeri okunamadı", "Takip oranı azalışı"])
            arg = ["--parametreler", str(p), "--senaryolar", str(s), "--cikti", str(Path(tmp) / "x.xlsx")]
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(arg + ["--kar-tamponu-yok"]), 0)
            self.assertEqual(main.main(arg + ["--asgari", "13", "--hedef", "12"]), 2)
            self.assertEqual(main.main(["--parametreler", str(Path(tmp) / "yok.csv")]), 1)
            p.write_text("Alan;Değer\nÖzkaynak;1.500\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)                                       # RAV ve krediler eksik


if __name__ == "__main__":
    unittest.main()
