"""
Stres Testi Senaryo Hesabı — Gerisi Bende kod bloğu
İç Kontrol ve Risk Yönetimi › Risk Yönetimi Uzmanı

Kur, faiz ve takip oranı şoklarının kârlılık ve sermaye yeterliliği üzerindeki etkisini senaryo bazında hesaplar.
Girdiler bankanın kendi duyarlılık rakamlarıdır; kod bunları senaryolara uygular (basit, doğrusal bir duyarlılık modeli).
  Kur etkisi          = net yabancı para pozisyonu × kur şoku %            (uzun pozisyon TL değer kaybında kazanır)
  Faiz etkisi         = (100 bp'lik net faiz geliri duyarlılığı + 100 bp'lik değerleme duyarlılığı) × şok bp ÷ 100
  Ek takip            = toplam krediler × takip oranı artışı (puan) ÷ 100
  Kredi zararı        = ek takip × karşılık oranı %
  Toplam şok etkisi   = kur etkisi + faiz etkisi − kredi zararı
  Stres sonrası kâr   = baz yıllık kâr (vergi öncesi) + toplam şok etkisi
  Sermaye etkisi      = min(0, stres sonrası kâr)      (zarar önce yılın kârından karşılanır; --kar-tamponu-yok ile
                        min(0, toplam şok etkisi)). Olumlu etki sermayeye eklenmez; vergi etkisi hesaplanmaz.
  Stres sonrası RAV   = RAV + yabancı para RAV × kur şoku %      (YP varlıkların TL karşılığı büyür; ihtiyatlı olmak
                        için karşılık ayrılan tutar RAV'dan düşülmez)
  Sermaye yeterliliği = (özkaynak + sermaye etkisi) ÷ stres sonrası RAV
  Sermaye ihtiyacı    = hedef oranı tutturmak için gereken ek özkaynak.
  Ters stres          : diğer şoklar sıfırken, tek başına hangi büyüklükte şok oranı asgariye indirir (ikiye bölme ile).
Asgari (%8) ve hedef (%12) oranlar parametredir; BDDK düzenlemelerindeki güncel oran ve tamponları kontrol edin.
Yasal stres testi / İSEDES raporlamasının yerini tutmaz. İnternete bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: 5 senaryo
    python main.py --parametreler banka.csv --senaryolar senaryolar.csv --asgari 8 --hedef 12
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


ALANLAR = {"ozkaynak": ("ozkaynak", "yasal ozkaynak"), "rav": ("risk agirlikli varliklar", "rav", "toplam rav"), "yp_rav": ("yabanci para rav", "yp rav"),
           "yp_poz": ("net yabanci para pozisyonu", "net yp pozisyonu", "yp net genel pozisyon"), "nfg": ("net faiz geliri duyarliligi", "nfg duyarliligi", "100 bp net faiz geliri etkisi"),
           "deger": ("degerleme duyarliligi", "100 bp degerleme etkisi", "ekonomik deger duyarliligi"), "kredi": ("toplam krediler", "krediler"),
           "takip": ("takipteki krediler", "takip", "donuk alacaklar"), "karsilik": ("karsilik orani", "takip karsilik orani"), "kar": ("yillik kar", "vergi oncesi kar", "baz yillik kar")}
SENARYO_SUTUNLARI = {"ad": ("senaryo", "senaryo adi"), "kur": ("kur soku", "kur soku yuzde", "kur"), "faiz": ("faiz soku", "faiz soku bp", "faiz"),
                     "takip": ("takip orani artisi", "takip artisi", "takip orani artisi puan"), "karsilik": ("karsilik orani", "karsilik orani yuzde"),
                     "aciklama": ("aciklama", "not")}
ZORUNLU = (("ozkaynak", "Özkaynak"), ("rav", "Risk Ağırlıklı Varlıklar"), ("kredi", "Toplam Krediler"))


def parametre_oku(yol: Path) -> tuple[dict, list[dict]]:
    ham = {katla(re.sub(r"\(.*?\)", "", metin(r[0]))): r[1] for r in tablo_oku(yol) if len(r) > 1 and metin(r[0]) and r[1] not in (None, "")}
    p, uy = {}, []
    for alan, adlar in ALANLAR.items():
        v = next((ham[a] for a in adlar if a in ham), None)
        p[alan] = para(v)
        if v is not None and p[alan] is None:
            raise ValueError(f"{yol.name}: '{adlar[0]}' sayı olarak okunamadı: '{metin(v)}'")
    eksik = [ad for alan, ad in ZORUNLU if p[alan] is None or p[alan] <= 0]
    if eksik:
        raise ValueError(f"{yol.name}: zorunlu alan eksik veya sıfır: {', '.join(eksik)}")
    for alan, ad, varsayilan in (("yp_rav", "Yabancı Para RAV", 0), ("yp_poz", "Net Yabancı Para Pozisyonu", 0), ("nfg", "Net Faiz Geliri Duyarlılığı", 0), ("deger", "Değerleme Duyarlılığı", 0),
                                 ("takip", "Takipteki Krediler", 0), ("karsilik", "Karşılık Oranı", 60), ("kar", "Yıllık Kâr", 0)):
        if p[alan] is None:
            p[alan] = Decimal(varsayilan)
            uy.append(uyari("Orta" if alan in ("yp_poz", "nfg", "karsilik", "kar") else "Bilgi", "Parametre girilmemiş", ad, f"Varsayılan {varsayilan} kullanıldı; ilgili şokun etkisi eksik hesaplanır."))
    if p["yp_rav"] > p["rav"]:
        raise ValueError(f"{yol.name}: Yabancı Para RAV toplam RAV'dan büyük olamaz")
    return p, uy


def senaryo_oku(yol: Path) -> tuple[list[dict], list[dict]]:
    s, uy = [], []
    for r in kayitlar(yol, SENARYO_SUTUNLARI, ("ad",)):
        if not metin(r["ad"]):
            continue
        d = {"ad": metin(r["ad"]), "aciklama": metin(r["aciklama"])}
        for alan in ("kur", "faiz", "takip", "karsilik"):
            d[alan] = para(r[alan])
            if metin(r[alan]) and d[alan] is None:
                uy.append(uyari("Yüksek", "Senaryo değeri okunamadı", d["ad"], f"{alan}: '{metin(r[alan])}'; 0 alındı."))
            if alan != "karsilik":
                d[alan] = d[alan] or Decimal(0)
        if d["takip"] < 0:
            uy.append(uyari("Orta", "Takip oranı azalışı", d["ad"], "Takip oranı artışı eksi girilmiş; kredi zararı 0 alındı."))
            d["takip"] = Decimal(0)
        s.append(d)
    return s, uy


def uygula(p: dict, kur: Decimal, faiz: Decimal, takip: Decimal, karsilik: Decimal | None = None, kar_tamponu: bool = True) -> dict:
    """Tek bir şok bileşimi için etkiler. Oranlar yüzde, faiz baz puan, takip puan."""
    ko = (karsilik if karsilik is not None else p["karsilik"]) / 100
    kur_etki = p["yp_poz"] * kur / 100
    nfg_etki, deger_etki = p["nfg"] * faiz / 100, p["deger"] * faiz / 100
    ek_takip = p["kredi"] * takip / 100
    kredi_zarari = ek_takip * ko
    sok = kur_etki + nfg_etki + deger_etki - kredi_zarari
    kar = p["kar"] + sok
    sermaye_etki = min(Decimal(0), kar if kar_tamponu else sok)
    rav = p["rav"] + p["yp_rav"] * kur / 100
    ozkaynak = p["ozkaynak"] + sermaye_etki
    syr = ozkaynak / rav * 100 if rav > 0 else Decimal(0)
    takip_orani = (p["takip"] + ek_takip) / p["kredi"] * 100
    return {"kur_etki": kur_etki, "nfg_etki": nfg_etki, "deger_etki": deger_etki, "ek_takip": ek_takip, "kredi_zarari": kredi_zarari, "sok": sok, "kar": kar,
            "sermaye_etki": sermaye_etki, "rav": rav, "ozkaynak": ozkaynak, "syr": syr, "takip_orani": takip_orani}


def ters_stres(p: dict, asgari: Decimal, kar_tamponu: bool = True) -> list[dict]:
    """Her şok türü için, tek başına SYR'yi asgariye indiren büyüklük (olumsuz yönde)."""
    sonuc = []
    for ad, birim, ust, f in (("Kur şoku", "%", Decimal(300), lambda x: uygula(p, x, 0, 0, kar_tamponu=kar_tamponu)),
                              ("Faiz şoku", "bp", Decimal(5000), lambda x: uygula(p, 0, x, 0, kar_tamponu=kar_tamponu)),
                              ("Takip oranı artışı", "puan", Decimal(100), lambda x: uygula(p, 0, 0, x, kar_tamponu=kar_tamponu))):
        sinir = {1: ust, -1: Decimal(95) if ad == "Kur şoku" else ust}      # kur %100'den fazla düşemez
        yon = 1 if ad == "Takip oranı artışı" else next((y for y in (1, -1) if f(y * sinir[y])["syr"] < f(Decimal(0))["syr"]), None)
        ust = sinir[yon or 1]
        if f(Decimal(0))["syr"] < asgari:
            sonuc.append({"ad": ad, "birim": birim, "deger": Decimal(0), "not": "Baz durumda oran zaten asgarinin altında"})
        elif yon is None or f(yon * ust)["syr"] >= asgari:
            sonuc.append({"ad": ad, "birim": birim, "deger": None, "not": f"±{ust:f} {birim} aralığında oranı düşürmüyor" if yon is None else
                          f"{yon * ust:+f} {birim} şokta bile oran asgarinin üzerinde"})
        else:
            a, b = Decimal(0), ust
            for _ in range(60):
                m = (a + b) / 2
                if f(yon * m)["syr"] >= asgari:
                    a = m
                else:
                    b = m
            sonuc.append({"ad": ad, "birim": birim, "deger": yuvarla(yon * b, "0.1"), "not": ""})
    return sonuc


def hesapla(p: dict, senaryolar: list[dict], asgari: Decimal = Decimal(8), hedef: Decimal = Decimal(12), kar_tamponu: bool = True) -> dict:
    uy = []
    baz = uygula(p, Decimal(0), Decimal(0), Decimal(0), kar_tamponu=kar_tamponu)
    if baz["syr"] < hedef:
        uy.append(uyari("Orta" if baz["syr"] >= asgari else "Yüksek", "Baz oran hedefin altında", "Baz durum", f"Sermaye yeterliliği %{yuvarla(baz['syr'])}; hedef %{hedef:f}, asgari %{asgari:f}."))
    sonuc = []
    for s in senaryolar:
        r = uygula(p, s["kur"], s["faiz"], s["takip"], s["karsilik"], kar_tamponu)
        r |= {"ad": s["ad"], "aciklama": s["aciklama"], "kur": s["kur"], "faiz": s["faiz"], "takip": s["takip"], "karsilik": s["karsilik"] if s["karsilik"] is not None else p["karsilik"],
              "syr_degisim": r["syr"] - baz["syr"], "ihtiyac": max(Decimal(0), hedef / 100 * r["rav"] - r["ozkaynak"]),
              "asgari_ihtiyac": max(Decimal(0), asgari / 100 * r["rav"] - r["ozkaynak"]),
              "durum": "Asgarinin altında" if r["syr"] < asgari else "Hedefin altında" if r["syr"] < hedef else "Yeterli"}
        etkiler = {"Kur": r["kur_etki"], "Faiz (net faiz geliri)": r["nfg_etki"], "Faiz (değerleme)": r["deger_etki"], "Kredi zararı": -r["kredi_zarari"]}
        r["en_buyuk"] = min(etkiler, key=etkiler.get) if min(etkiler.values()) < 0 else "—"
        if r["durum"] != "Yeterli":
            uy.append(uyari("Yüksek" if r["durum"] == "Asgarinin altında" else "Orta", r["durum"], s["ad"],
                            f"Sermaye yeterliliği %{yuvarla(r['syr'])} (baz %{yuvarla(baz['syr'])}). En büyük olumsuz etki: {r['en_buyuk']}. Hedef için ek özkaynak ihtiyacı {r['ihtiyac']:,.0f}."
                            .replace(",", ".")))
        if r["kar"] < 0 and r["durum"] == "Yeterli":
            uy.append(uyari("Bilgi", "Senaryoda dönem zararı", s["ad"], f"Stres sonrası kâr {r['kar']:,.0f}; oran hedefin üzerinde kalıyor.".replace(",", ".")))
        sonuc.append(r)
    return {"p": p, "baz": baz, "senaryolar": sonuc, "ters": ters_stres(p, asgari, kar_tamponu), "asgari": asgari, "hedef": hedef, "kar_tamponu": kar_tamponu, "uyarilar": uy}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

DURUM_RENK = {"Yeterli": YESIL, "Hedefin altında": SARI, "Asgarinin altında": KIRMIZI}
ALAN_ADI = {"ozkaynak": "Özkaynak", "rav": "Risk ağırlıklı varlıklar", "yp_rav": "Yabancı para RAV", "yp_poz": "Net yabancı para pozisyonu (+ uzun)",
            "nfg": "Net faiz geliri duyarlılığı (+100 bp, 12 ay)", "deger": "Değerleme duyarlılığı (+100 bp)", "kredi": "Toplam krediler", "takip": "Takipteki krediler",
            "karsilik": "Karşılık oranı %", "kar": "Baz yıllık kâr (vergi öncesi)"}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb, baz = Workbook(), r["baz"]
    sn = wb.active
    sn.title = "Senaryo Sonuçları"
    sn.append([f"Baz sermaye yeterliliği %{yuvarla(baz['syr'])} · asgari %{r['asgari']:f} · hedef %{r['hedef']:f} · " +
               ("zarar önce yılın kârından karşılanır" if r["kar_tamponu"] else "kâr tamponu kullanılmaz")])
    sn["A1"].font = Font(bold=True)
    _baslik(sn, ["Senaryo", "Kur Şoku %", "Faiz Şoku (bp)", "Takip Artışı (puan)", "Toplam Şok Etkisi", "Stres Sonrası Kâr", "Sermaye Etkisi", "Stres Sonrası Özkaynak",
                 "Stres Sonrası RAV", "Sermaye Yeterliliği %", "Değişim (puan)", "Takip Oranı %", "Durum", "Hedef İçin Ek Özkaynak", "En Büyük Olumsuz Etki", "Açıklama"],
            (30, 9, 9, 9, 20, 20, 20, 20, 20, 11, 9, 9, 16, 20, 22, 50), "B3")
    _satir(sn, ["Baz durum", 0, 0, 0, 0, baz["kar"], 0, baz["ozkaynak"], baz["rav"], yuvarla(baz["syr"]), 0, yuvarla(baz["takip_orani"]),
                "Yeterli" if baz["syr"] >= r["hedef"] else "Hedefin altında" if baz["syr"] >= r["asgari"] else "Asgarinin altında", "", "", ""], para_s=(5, 6, 7, 8, 9))
    for s in r["senaryolar"]:
        _satir(sn, [s["ad"], s["kur"], s["faiz"], s["takip"], yuvarla(s["sok"]), yuvarla(s["kar"]), yuvarla(s["sermaye_etki"]), yuvarla(s["ozkaynak"]), yuvarla(s["rav"]),
                    yuvarla(s["syr"]), yuvarla(s["syr_degisim"]), yuvarla(s["takip_orani"]), s["durum"], yuvarla(s["ihtiyac"]), s["en_buyuk"], s["aciklama"]],
               para_s=(5, 6, 7, 8, 9, 14), dolgu={13: DURUM_RENK[s["durum"]], 6: KIRMIZI if s["kar"] < 0 else None}, sar=(16,))

    et = wb.create_sheet("Etki Dökümü")
    _baslik(et, ["Senaryo", "Kur Etkisi", "Faiz: Net Faiz Geliri", "Faiz: Değerleme", "Ek Takip", "Karşılık Oranı %", "Kredi Zararı", "Toplam Şok Etkisi", "RAV Değişimi"],
            (30, 20, 20, 20, 20, 10, 20, 20, 20))
    for s in r["senaryolar"]:
        _satir(et, [s["ad"], yuvarla(s["kur_etki"]), yuvarla(s["nfg_etki"]), yuvarla(s["deger_etki"]), yuvarla(s["ek_takip"]), s["karsilik"], yuvarla(-s["kredi_zarari"]), yuvarla(s["sok"]),
                    yuvarla(s["rav"] - r["p"]["rav"])], para_s=(2, 3, 4, 5, 7, 8, 9),
               dolgu={j: KIRMIZI for j, v in ((2, s["kur_etki"]), (3, s["nfg_etki"]), (4, s["deger_etki"]), (7, -s["kredi_zarari"])) if v < 0})
    et.append([])
    et.append(["Eksi tutar zarar, artı tutar kazançtır. Etkiler doğrusal duyarlılıkla hesaplanır; şoklar arasındaki etkileşim ve ikinci tur etkiler modellenmez."])

    ts = wb.create_sheet("Ters Stres")
    _baslik(ts, ["Şok Türü", f"Oranı %{r['asgari']:f} Düzeyine İndiren Şok", "Birim", "Not"], (24, 26, 8, 70))
    for t in r["ters"]:
        _satir(ts, [t["ad"], t["deger"], t["birim"], t["not"]])
    ts.append([])
    ts.append(["Her satırda yalnız o şok uygulanır, diğerleri sıfırdır. Değer küçükse banka o risk türüne karşı kırılgandır."])

    pr = wb.create_sheet("Parametreler")
    _baslik(pr, ["Parametre", "Değer"], (48, 22))
    for alan, ad in ALAN_ADI.items():
        _satir(pr, [ad, r["p"][alan]], para_s=(2,))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(p_yol: Path, s_yol: Path, cikti: Path, asgari: Decimal = Decimal(8), hedef: Decimal = Decimal(12), kar_tamponu: bool = True) -> dict:
    p, uy = parametre_oku(p_yol)
    senaryolar, uy2 = senaryo_oku(s_yol)
    if not senaryolar:
        raise ValueError(f"{s_yol.name}: senaryo bulunamadı")
    r = hesapla(p, senaryolar, asgari, hedef, kar_tamponu)
    r["uyarilar"] = sirala(uy + uy2 + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Kur, faiz ve takip oranı şoklarının kârlılık ve sermaye yeterliliğine etkisini senaryo bazında hesaplar.")
    p.add_argument("--parametreler", type=Path, default=ORNEK / "parametreler.csv", help="Alan;Değer: Özkaynak, Risk Ağırlıklı Varlıklar, Net Yabancı Para Pozisyonu, duyarlılıklar...")
    p.add_argument("--senaryolar", type=Path, default=ORNEK / "senaryolar.csv", help="Senaryo, Kur Şoku (%%), Faiz Şoku (bp), Takip Oranı Artışı (puan), Karşılık Oranı (%%), Açıklama")
    p.add_argument("--asgari", type=Decimal, default=Decimal(8), help="Asgari sermaye yeterliliği oranı %% (varsayılan 8)")
    p.add_argument("--hedef", type=Decimal, default=Decimal(12), help="Hedef oran %% (varsayılan 12)")
    p.add_argument("--kar-tamponu-yok", action="store_true", help="Zararı yılın kârından karşılama; şok etkisini doğrudan sermayeden düş")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "stres_testi.xlsx")
    a = p.parse_args(argv)
    for y in (a.parametreler, a.senaryolar):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    if not (0 < a.asgari <= a.hedef):
        print("[X] 0 < --asgari ≤ --hedef olmalı")
        return 2
    try:
        r = calistir(a.parametreler, a.senaryolar, a.cikti, a.asgari, a.hedef, not a.kar_tamponu_yok)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] Baz sermaye yeterliliği %{yuvarla(r['baz']['syr'])} · {len(r['senaryolar'])} senaryo")
    for s in r["senaryolar"]:
        print(f"     {s['ad']}: %{yuvarla(s['syr'])} ({yuvarla(s['syr_degisim']):+} puan) · {s['durum']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
