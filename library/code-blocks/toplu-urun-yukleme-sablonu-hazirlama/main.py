"""
Toplu Ürün Yükleme Şablonu Hazırlama — Workers / Workless kod bloğu
E-ticaret › Pazaryeri Yönetimi › Pazaryeri Uzmanı

Ürün ana veri dosyasını pazaryerinin toplu ürün yükleme şablonuna dönüştürür. Pazaryerlerinin şablonları
farklıdır ve zamanla değişir; bu yüzden dönüşüm bir eşleştirme dosyasıyla (JSON) tanımlanır:
  - Şablon sütunu ← ana veri sütunu, sabit değer veya birleşik metin ("{Marka} {Ürün Adı} {Renk}").
  - Kategori, renk gibi değerler için eşleştirme tablosu; karşılığı olmayan değerler ayrı listelenir.
  - Kontroller: zorunlu alan, uzunluk sınırı, izinli değerler (ör. KDV 0/1/10/20), sayı/para biçimi, barkod
    (GTIN-8/12/13/14 kontrol hanesi), benzersiz barkod ve stok kodu, görsel adresi (https), satış fiyatının
    liste fiyatını aşması, varyant tutarlılığı (aynı modelde marka/kategori aynı, renk-beden tekrarı yok).
  - Hatasız ürünler şablonun sütun sırasıyla yükleme dosyasına yazılır (şablon .xlsx ise kopyası doldurulur);
    hatalı ürünler ve eşleşmeyen değerler rapora yazılır.
İnternete bağlanmaz; pazaryerine hiçbir şey yüklemez.

Kullanım:
    python main.py                                                    # örnek: kurgusal pazaryeri şablonu
    python main.py --girdi urunler.xlsx --sablon pazaryeri_sablonu.xlsx --eslestirme eslestirme.json
    python main.py --girdi urunler.xlsx --sablon sablon.xlsx --eslestirme eslestirme.json --hepsi
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import shutil
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
GORSEL_UZANTI = (".jpg", ".jpeg", ".png", ".webp")


def katla(s) -> str:
    return " ".join(str(s or "").translate(_TR).lower().split())


def metin(x) -> str:
    if x is None:
        return ""
    if isinstance(x, float) and x.is_integer():
        return str(int(x))
    return str(x).strip()


def sayi(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1:
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return satirlar


def gtin_kontrol(kod: str) -> str:
    """GTIN-8/12/13/14 (EAN/UPC): sağdan sola 3-1 ağırlıklı toplamla kontrol hanesi. Boş metin = geçerli."""
    if not kod.isdigit():
        return "Barkod yalnız rakamdan oluşmalı"
    if len(kod) not in (8, 12, 13, 14):
        return f"Barkod uzunluğu {len(kod)}; GTIN 8, 12, 13 veya 14 hane olmalı"
    govde = kod[:-1]
    toplam = sum(int(c) * (3 if i % 2 == 0 else 1) for i, c in enumerate(reversed(govde)))
    return "" if (10 - toplam % 10) % 10 == int(kod[-1]) else "Barkod kontrol hanesi tutmuyor (yazım hatası olabilir)"


# ----------------------------------------------------------------------------
# Ayarlar ve girdiler
# ----------------------------------------------------------------------------

def eslestirme_oku(yol: Path) -> tuple[dict, dict[str, dict[str, dict]]]:
    """Eşleştirme JSON'u ve eşleştirme tabloları: ad → {katlanmış anahtar: {sütun: değer}}"""
    ayar = json.loads(yol.read_text(encoding="utf-8"))
    if not ayar.get("alanlar"):
        raise SystemExit(f"{yol.name}: 'alanlar' tanımı yok")
    tablolar = {}
    for ad, e in ayar.get("eslestirmeler", {}).items():
        if "degerler" in e:
            tablolar[ad] = {katla(k): {"deger": v} for k, v in e["degerler"].items()}
        else:
            s = [r for r in tablo_oku(yol.parent / e["dosya"]) if any(c not in (None, "") for c in r)]
            bas = [metin(x) for x in s[0]]
            if e["anahtar"] not in bas:
                raise SystemExit(f"{e['dosya']}: '{e['anahtar']}' sütunu yok. Başlıklar: {bas}")
            ai = bas.index(e["anahtar"])
            tablolar[ad] = {katla(r[ai]): {b: metin(r[i]) if i < len(r) else "" for i, b in enumerate(bas)} for r in s[1:] if metin(r[ai])}
    return ayar, tablolar


def sablon_sutunlari(yol: Path, satir: int) -> list[str]:
    s = tablo_oku(yol)
    if len(s) < satir:
        raise SystemExit(f"{yol.name}: {satir}. satırda başlık yok")
    return [metin(x) for x in s[satir - 1] if metin(x)]


def urunleri_oku(yol: Path, kaynaklar: set[str]) -> list[dict]:
    s = [r for r in tablo_oku(yol) if any(c not in (None, "") for c in r)]
    istenen = {katla(k) for k in kaynaklar}
    bi = max(range(min(10, len(s))), key=lambda i: sum(katla(x) in istenen for x in s[i]))   # başlık satırı
    bas = [metin(x) for x in s[bi]]
    return [{"_satir": bi + 2 + j, **{b: r[i] if i < len(r) else None for i, b in enumerate(bas) if b}} for j, r in enumerate(s[bi + 1:])]


# ----------------------------------------------------------------------------
# Dönüştürme ve kontrol
# ----------------------------------------------------------------------------

@dataclass
class Urun:
    kaynak: dict
    cikti: dict = field(default_factory=dict)
    bulgular: list[tuple[str, str, str]] = field(default_factory=list)

    @property
    def hatali(self) -> bool:
        return any(b[0] == "Hata" for b in self.bulgular)

    def deger(self, ad: str) -> str:
        return metin(self.cikti.get(ad)) or metin(self.kaynak.get(ad))


def _kaynak(urun: dict, ad: str):
    if ad in urun:
        return urun[ad]
    k = katla(ad)
    return next((v for a, v in urun.items() if katla(a) == k), None)


def donustur(urunler: list[dict], ayar: dict, tablolar: dict, sutunlar: list[str]) -> tuple[list[Urun], Counter]:
    alanlar = ayar["alanlar"]
    eslesmeyen, sayilan = Counter(), set()
    sonuc = []
    sira = [a for a in sutunlar if a in alanlar and "sablon" not in alanlar[a]] + [a for a in sutunlar if a in alanlar and "sablon" in alanlar[a]]
    for ham in urunler:
        u = Urun(ham)
        for ad in sira:
            o = alanlar[ad]
            if "sabit" in o:
                u.cikti[ad] = o["sabit"]
                continue
            if "sablon" in o:
                d = re.sub(r"\{([^}]+)\}", lambda m: u.deger(m.group(1)), o["sablon"])
                d = " ".join(d.split())
            else:
                d = metin(_kaynak(ham, o["kaynak"]))
                if "bol" in o:
                    parca = [p.strip() for p in d.split(o["bol"]) if p.strip()]
                    d = parca[o.get("sira", 1) - 1] if len(parca) >= o.get("sira", 1) else ""
            if not d:
                if o.get("zorunlu"):
                    u.bulgular.append(("Hata", ad, "Zorunlu alan boş"))
                u.cikti[ad] = ""
                continue
            if "eslestirme" in o:
                kayit = tablolar.get(o["eslestirme"], {}).get(katla(d))
                if kayit is None:
                    if (o["eslestirme"], d, id(ham)) not in sayilan:      # aynı ürün iki sütunda aynı eşleştirmeyi kullanabilir
                        sayilan.add((o["eslestirme"], d, id(ham)))
                        eslesmeyen[(o["eslestirme"], d)] += 1
                    u.bulgular.append(("Hata", ad, f"'{d}' için {o['eslestirme']} eşleştirmesi yok"))
                    u.cikti[ad] = ""
                    continue
                d = kayit.get(o.get("sutun", "deger"), "")
            tur = o.get("tur")
            if tur in ("para", "sayi", "tamsayi"):
                n = sayi(d)
                if n is None:
                    u.bulgular.append(("Hata", ad, f"Sayı değil: '{d}'"))
                elif tur == "tamsayi" and n != n.to_integral_value():
                    u.bulgular.append(("Hata", ad, f"Tam sayı olmalı: {d}"))
                elif n < 0 or (tur == "para" and n == 0):
                    u.bulgular.append(("Hata", ad, f"Geçersiz değer: {d}"))
                d = n if n is None or tur != "tamsayi" else int(n)
            if "izinli" in o:
                izin = {katla(x) for x in o["izinli"]}
                if katla(d) not in izin:
                    u.bulgular.append(("Hata", ad, f"'{metin(d)}' izinli değil; izinli: {', '.join(map(str, o['izinli']))}"))
            if isinstance(d, str):
                if o.get("en_fazla") and len(d) > o["en_fazla"]:
                    u.bulgular.append(("Hata", ad, f"{len(d)} karakter; en fazla {o['en_fazla']}"))
                if o.get("en_az") and len(d) < o["en_az"]:
                    u.bulgular.append(("Hata", ad, f"{len(d)} karakter; en az {o['en_az']}"))
            if o.get("kontrol") == "gtin":
                h = gtin_kontrol(metin(d))
                if h:
                    u.bulgular.append(("Hata", ad, h))
            elif o.get("kontrol") == "url":
                if not re.match(r"https://\S+$", metin(d)):
                    u.bulgular.append(("Hata", ad, "Görsel adresi https:// ile başlamalı ve boşluk içermemeli"))
                elif not metin(d).lower().split("?")[0].endswith(GORSEL_UZANTI):
                    u.bulgular.append(("Dikkat", ad, "Görsel uzantısı jpg/png/webp değil; pazaryerinin kabul ettiği biçimi kontrol edin"))
            u.cikti[ad] = d
        sonuc.append(u)
    return sonuc, eslesmeyen


def capraz_kontroller(urunler: list[Urun], ayar: dict) -> None:
    alanlar = ayar["alanlar"]
    for ad, o in alanlar.items():
        if o.get("benzersiz"):
            ilk = {}
            for u in urunler:
                d = metin(u.cikti.get(ad))
                if d and d in ilk:
                    u.bulgular.append(("Hata", ad, f"'{d}' tekrar ediyor (kaynak satır {ilk[d]})"))
                elif d:
                    ilk[d] = u.kaynak["_satir"]
    k = ayar.get("kurallar", {})
    if "fiyat" in k:
        s, l_ = k["fiyat"]["satis"], k["fiyat"]["liste"]
        for u in urunler:
            a, b = u.cikti.get(s), u.cikti.get(l_)
            if isinstance(a, Decimal) and isinstance(b, Decimal) and a > b:
                u.bulgular.append(("Hata", s, f"Satış fiyatı ({str(a).replace('.', ',')}) liste/piyasa fiyatından ({str(b).replace('.', ',')}) yüksek"))
    if "varyant" in k:
        v = k["varyant"]
        gruplar = defaultdict(list)
        for u in urunler:
            if metin(u.cikti.get(v["model"])):
                gruplar[metin(u.cikti[v["model"]])].append(u)
        for model, us in gruplar.items():
            for ad in v.get("ortak", []):
                ilk = us[0].cikti.get(ad)
                for u in us[1:]:
                    if u.cikti.get(ad) != ilk and metin(u.cikti.get(ad)) and metin(ilk):
                        u.bulgular.append(("Hata", ad, f"Aynı modelin ({model}) varyantlarında farklı: '{metin(u.cikti.get(ad))}' / '{metin(ilk)}'"))
            gorulen = {}
            for u in us:
                anahtar = tuple(katla(u.cikti.get(a)) for a in v.get("ayirt", []))
                if anahtar in gorulen:
                    u.bulgular.append(("Hata", "Varyant", f"{model} modelinde {' / '.join(metin(u.cikti.get(a)) for a in v['ayirt'])} varyantı "
                                                         f"tekrar ediyor (kaynak satır {gorulen[anahtar]})"))
                else:
                    gorulen[anahtar] = u.kaynak["_satir"]


# ----------------------------------------------------------------------------
# Çıktılar
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Hata": "FDE2E1", "Dikkat": "FFF4CE", "Bilgi": "E8F0FE"}


def _hucre(d):
    if isinstance(d, Decimal):
        return float(d)
    return d


def yukleme_yaz(cikti: Path, sablon: Path, satir: int, sutunlar: list[str], urunler: list[Urun], ayar: dict) -> int:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    if sablon.suffix.lower() in {".xlsx", ".xlsm"}:
        shutil.copyfile(sablon, cikti)
        wb = load_workbook(cikti)
        ws = wb.worksheets[0]
        konum = {metin(c.value): c.column for c in ws[satir] if metin(c.value)}
        bas = ayar.get("veri_baslangic_satiri", satir + 1)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Ürünler"
        ws.append(sutunlar)
        for h in ws[1]:
            h.font = Font(bold=True)
        konum = {a: i for i, a in enumerate(sutunlar, 1)}
        bas = 2
    metin_sutunlari = {a for a, o in ayar["alanlar"].items() if o.get("kontrol") == "gtin" or a in ("Model Kodu", "Stok Kodu")}
    for i, u in enumerate(urunler):
        for ad, c in konum.items():
            if ad in u.cikti:
                h = ws.cell(bas + i, c, _hucre(u.cikti[ad]) if ad not in metin_sutunlari else metin(u.cikti[ad]))
                if ad in metin_sutunlari:
                    h.number_format = "@"                       # barkodun başındaki sıfırlar kaybolmasın
    wb.save(cikti)
    return len(urunler)


def rapor_yaz(cikti: Path, urunler: list[Urun], eslesmeyen: Counter, ayar_bulgulari: list, ayar: dict, yuklenen: int) -> None:
    wb = Workbook()
    o = wb.active
    o.title = "Özet"
    o.column_dimensions["A"].width, o.column_dimensions["B"].width = 34, 70
    hatali = [u for u in urunler if u.hatali]
    for etiket, deger in [("Pazaryeri", ayar.get("pazaryeri", "")), ("Ürün (varyant) satırı", len(urunler)),
                          ("Yükleme dosyasına yazılan", yuklenen), ("Hatalı ürün", len(hatali)),
                          ("Yazılmayan (hatalı)", len(urunler) - yuklenen), ("Eşleşmeyen değer", len(eslesmeyen))]:
        o.append([etiket, deger])
        o.cell(o.max_row, 1).font = Font(bold=True)
    o.append([])
    o.append(["En sık sorunlar", "Adet"])
    for h in o[o.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    sayac = Counter(f"{a}: {re.sub(chr(39) + '.*?' + chr(39), '…', m)}" for u in urunler for on, a, m in u.bulgular if on == "Hata")
    for s, n in sayac.most_common(10):
        o.append([s, n])
    o.append([])
    o.append(["Not", "Kurallar eşleştirme dosyasından gelir. Pazaryerinin güncel şablonu, zorunlu alanları, kategori özellikleri ve "
                     "görsel kuralları satıcı panelinden kontrol edilmelidir. Yükleme sonrası pazaryerinin hata raporunu da inceleyin."])
    o.cell(o.max_row, 2).alignment = Alignment(wrap_text=True, vertical="top")

    h = wb.create_sheet("Hatalar")
    bas = ["Önem", "Kaynak Satır", "Model", "Stok Kodu", "Barkod", "Alan", "Sorun", "Düzeltildi mi?"]
    h.append(bas)
    for c in h[1]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI
    for j, w in enumerate((9, 11, 12, 18, 16, 28, 80, 14), 1):
        h.column_dimensions[get_column_letter(j)].width = w
    model_ad = ayar.get("kurallar", {}).get("varyant", {}).get("model", "Model Kodu")
    barkod_ad = next((a for a, x in ayar["alanlar"].items() if x.get("kontrol") == "gtin"), "Barkod")
    for u in urunler:
        for on, a, m in u.bulgular:
            h.append([on, u.kaynak["_satir"], u.deger(model_ad), metin(_kaynak(u.kaynak, "Stok Kodu")), u.deger(barkod_ad), a, m, ""])
            h.cell(h.max_row, 1).fill = PatternFill("solid", fgColor=RENK[on])
            h.cell(h.max_row, 8).fill = PatternFill("solid", fgColor="FFF4CE")
    h.auto_filter.ref = h.dimensions
    h.freeze_panes = "A2"

    e = wb.create_sheet("Eşleşmeyen Değerler")
    e.append(["Eşleştirme", "Bizim Değer", "Ürün Sayısı", "Pazaryeri Karşılığı (doldurun)"])
    for c in e[1]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI
    for j, w in enumerate((14, 40, 12, 36), 1):
        e.column_dimensions[get_column_letter(j)].width = w
    for (ad, d), n in sorted(eslesmeyen.items()):
        e.append([ad, d, n, ""])
        e.cell(e.max_row, 4).fill = PatternFill("solid", fgColor="FFF4CE")

    k = wb.create_sheet("Ayar Kontrolü")
    k.append(["Önem", "Konu"])
    for c in k[1]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI
    k.column_dimensions["B"].width = 110
    for on, m in ayar_bulgulari or [("Bilgi", "Şablon ve eşleştirme sütunları uyumlu")]:
        k.append([on, m])
        k.cell(k.max_row, 1).fill = PatternFill("solid", fgColor=RENK[on])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


# ----------------------------------------------------------------------------
# Akış
# ----------------------------------------------------------------------------

def calistir(girdi: Path, sablon: Path, eslestirme: Path, cikti_klasoru: Path, hepsi: bool = False) -> dict:
    ayar, tablolar = eslestirme_oku(eslestirme)
    satir = ayar.get("sablon_baslik_satiri", 1)
    sutunlar = sablon_sutunlari(sablon, satir)
    ayar_bulgulari = [("Dikkat", f"Eşleştirmede tanımlı '{a}' şablonda yok; yazılmayacak") for a in ayar["alanlar"] if a not in sutunlar]
    bos = [a for a in sutunlar if a not in ayar["alanlar"]]
    if bos:
        ayar_bulgulari.append(("Bilgi", "Şablonda olup eşleştirmede tanımı olmayan sütunlar boş kalacak: " + ", ".join(bos)))
    kaynaklar = {o["kaynak"] for o in ayar["alanlar"].values() if "kaynak" in o}
    kaynaklar |= {m for o in ayar["alanlar"].values() for m in re.findall(r"\{([^}]+)\}", o.get("sablon", ""))}
    ham = urunleri_oku(girdi, kaynaklar)
    if ham:
        eksik = sorted(k for k in kaynaklar if _kaynak(ham[0], k) is None and k not in ayar["alanlar"])
        if eksik:
            ayar_bulgulari.append(("Hata", "Ana veride bulunamayan sütunlar: " + ", ".join(eksik)))
    urunler, eslesmeyen = donustur(ham, ayar, tablolar, sutunlar)
    capraz_kontroller(urunler, ayar)
    yazilacak = urunler if hepsi else [u for u in urunler if not u.hatali]
    yukleme = cikti_klasoru / f"yukleme_{re.sub(r'[^a-z0-9]+', '_', katla(ayar.get('pazaryeri', 'pazaryeri'))).strip('_')}.xlsx"
    n = yukleme_yaz(yukleme, sablon, satir, sutunlar, yazilacak, ayar)
    rapor = cikti_klasoru / "yukleme_kontrol_raporu.xlsx"
    rapor_yaz(rapor, urunler, eslesmeyen, ayar_bulgulari, ayar, n)
    return {"urunler": urunler, "eslesmeyen": eslesmeyen, "ayar_bulgulari": ayar_bulgulari, "yukleme": yukleme, "rapor": rapor, "yazilan": n}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Ürün ana verisini pazaryeri toplu yükleme şablonuna dönüştürür ve kontrol eder.")
    p.add_argument("--girdi", type=Path, default=ORNEK / "urun_ana_veri.csv", help="Ürün ana verisi (.xlsx/.csv), her satır bir varyant")
    p.add_argument("--sablon", type=Path, default=ORNEK / "pazaryeri_sablonu.csv", help="Pazaryeri şablonu (.xlsx önerilir, ya da başlık satırlı .csv)")
    p.add_argument("--eslestirme", type=Path, default=ORNEK / "eslestirme.json", help="Sütun eşleştirme ve kurallar (JSON)")
    p.add_argument("--cikti", type=Path, default=Path("cikti"), help="Çıktı klasörü")
    p.add_argument("--hepsi", action="store_true", help="Hatalı ürünleri de yükleme dosyasına yaz")
    a = p.parse_args(argv)
    s = calistir(a.girdi, a.sablon, a.eslestirme, a.cikti, a.hepsi)
    hatali = sum(1 for u in s["urunler"] if u.hatali)
    print(f"[OK] {len(s['urunler'])} ürün satırı · {s['yazilan']} yükleme dosyasına yazıldı · {hatali} hatalı")
    for on, m in s["ayar_bulgulari"]:
        print(f"[{'X' if on == 'Hata' else '!' if on == 'Dikkat' else 'i'}] {m}")
    for (ad, d), n in s["eslesmeyen"].most_common(5):
        print(f"[!] Eşleşmeyen {ad}: '{d}' ({n} ürün)")
    print(f"[OK] Yükleme dosyası: {s['yukleme'].resolve()}")
    print(f"[OK] Kontrol raporu: {s['rapor'].resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
