import json
import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))

import main  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"


def calistir(tmp, **kw):
    return main.calistir(ORNEK / "urun_ana_veri.csv", kw.pop("sablon", ORNEK / "pazaryeri_sablonu.csv"), ORNEK / "eslestirme.json", Path(tmp), **kw)


class OrnekTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        with tempfile.TemporaryDirectory() as tmp:
            cls.s = calistir(tmp)
            cls.yukleme = [list(r) for r in load_workbook(cls.s["yukleme"]).active.iter_rows(values_only=True)]
            cls.rapor = load_workbook(cls.s["rapor"])
        cls.u = {u.kaynak["Stok Kodu"]: u for u in cls.s["urunler"]}

    def bul(self, kod):
        return "\n".join(f"{a}: {m}" for _, a, m in self.u[kod].bulgular)

    def test_donusum(self):
        u = self.u["TS-100-NVY-S"]
        self.assertEqual(u.cikti["Ürün Adı"], "Kurgu Wear Basic Oversize Tişört Lacivert")      # Navy → Lacivert, başlıkta da
        self.assertEqual((u.cikti["Kategori"], u.cikti["Kategori ID"], u.cikti["Para Birimi"]), ("Tişört", "1001", "TRY"))
        self.assertEqual(u.cikti["Görsel 2"], "https://cdn.ornek-magaza.com/urun/ts100-navy-2.jpg")
        self.assertEqual(u.cikti["Görsel 3"], "")
        self.assertEqual((u.cikti["Satış Fiyatı (KDV Dahil)"], u.cikti["KDV Oranı"], u.cikti["Ürün Stok Adedi"]), (D("449.90"), 10, 12))
        self.assertEqual(self.u["SW-220-GRI-M"].cikti["Renk"], "Gri")

    def test_hatalar(self):
        self.assertEqual({k for k, u in self.u.items() if not u.hatali}, {"TS-100-SYH-S", "TS-100-SYH-M", "TS-100-SYH-L", "TS-100-NVY-S", "SW-220-GRI-M"})
        self.assertIn("Satış fiyatı (649,90) liste/piyasa fiyatından (599,90) yüksek", self.bul("TS-100-NVY-M"))
        self.assertIn("Barkod: '8690000120055' tekrar ediyor (kaynak satır 6)", self.bul("TS-100-NVY-L"))
        self.assertIn("Görsel 1: Zorunlu alan boş", self.bul("SW-220-GRI-L"))
        self.assertIn("Gri / M varyantı tekrar ediyor (kaynak satır 8)", self.bul("SW-220-GRI-M2"))
        self.assertIn("Barkod kontrol hanesi tutmuyor", self.bul("TR-300-SYH-36"))
        self.assertIn("KDV Oranı: '18' izinli değil", self.bul("TR-300-SYH-36"))
        self.assertIn("Görsel 2: Görsel adresi https://", self.bul("TR-300-SYH-38"))
        self.assertIn("Ürün Adı: 115 karakter; en fazla 100", self.bul("CN-410-BRD-STD"))
        self.assertIn("Renk: 'Bordo' için renk eşleştirmesi yok", self.bul("CN-410-BRD-STD"))
        self.assertIn("Ürün Açıklaması: 4 karakter; en az 30", self.bul("CN-411-SYH-STD"))
        self.assertEqual(dict(self.s["eslesmeyen"]), {("kategori", "Kadın > Aksesuar > Çanta"): 2, ("renk", "Bordo"): 1})

    def test_yukleme_dosyasi(self):
        self.assertEqual(self.s["yazilan"], 5)
        self.assertEqual(self.yukleme[0][:3], ["Barkod", "Model Kodu", "Marka"])
        self.assertEqual(len(self.yukleme), 6)
        self.assertEqual(self.yukleme[1][0], "8690000120017")                   # barkod metin olarak
        self.assertEqual(self.yukleme[1][8], 599.9)

    def test_rapor(self):
        self.assertEqual(self.rapor.sheetnames, ["Özet", "Hatalar", "Eşleşmeyen Değerler", "Ayar Kontrolü"])
        e = [list(r) for r in self.rapor["Eşleşmeyen Değerler"].iter_rows(min_row=2, values_only=True)]
        self.assertEqual(e, [["kategori", "Kadın > Aksesuar > Çanta", 2, None], ["renk", "Bordo", 1, None]])


class GtinTesti(unittest.TestCase):
    def test_gtin(self):
        self.assertEqual(main.gtin_kontrol("8690000120017"), "")
        self.assertEqual(main.gtin_kontrol("4006381333931"), "")                # bilinen geçerli EAN-13
        self.assertEqual(main.gtin_kontrol("96385074"), "")                     # geçerli EAN-8
        self.assertEqual(main.gtin_kontrol("036000291452"), "")                 # geçerli UPC-A
        self.assertIn("kontrol hanesi", main.gtin_kontrol("8690000120018"))
        self.assertIn("uzunluğu", main.gtin_kontrol("869000012001"[:-1]))
        self.assertIn("rakam", main.gtin_kontrol("86900A0120017"))


class XlsxSablonTesti(unittest.TestCase):
    def test_xlsx_sablon_kopyalanir(self):
        with tempfile.TemporaryDirectory() as tmp:
            wb = Workbook()
            ws = wb.active
            bas = main.sablon_sutunlari(ORNEK / "pazaryeri_sablonu.csv", 1)
            ws.append(bas)
            ws.append(["Örnek satır"] + [""] * (len(bas) - 1))
            wb.create_sheet("Talimatlar")["A1"] = "Pazaryeri talimatları"
            sablon = Path(tmp) / "sablon.xlsx"
            wb.save(sablon)
            ayar = json.loads((ORNEK / "eslestirme.json").read_text(encoding="utf-8"))
            ayar["veri_baslangic_satiri"] = 3
            (Path(tmp) / "kategori_eslestirme.csv").write_bytes((ORNEK / "kategori_eslestirme.csv").read_bytes())
            (Path(tmp) / "e.json").write_text(json.dumps(ayar, ensure_ascii=False), encoding="utf-8")
            s = main.calistir(ORNEK / "urun_ana_veri.csv", sablon, Path(tmp) / "e.json", Path(tmp) / "c")
            sonuc = load_workbook(s["yukleme"])
        self.assertEqual(sonuc.sheetnames, ["Sheet", "Talimatlar"])
        self.assertEqual(sonuc.active["A2"].value, "Örnek satır")                # açıklama satırı korunur
        self.assertEqual(sonuc.active["A3"].value, "8690000120017")
        self.assertEqual(sonuc.active.max_row, 7)


if __name__ == "__main__":
    unittest.main()
