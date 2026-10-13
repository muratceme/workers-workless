"""
İcra Dosyası Faiz ve Alacak Hesabı — Workers / Workless kod bloğu
İcra Takip › İcra Takip Elemanı (Avukat Katibi)

Asıl alacak, faiz başlangıcı ve faiz türüne göre dönemsel oranlarla işlemiş faizi, kısmi ödemelerin mahsubunu ve güncel borcu hesaplar.
  Dönem faizi        = asıl alacak × yıllık oran % ÷ 100 × gün sayısı ÷ --gun-esasi (365)         (basit faiz; faize faiz yürütülmez)
  Gün sayısı         = dönem bitişi − dönem başlangıcı                                          (başlangıç günü dahil, bitiş günü hariç)
  Oran değişimi      = oranlar dosyasındaki her "başlangıç tarihi" yeni bir dönem açar; tarih için o tarihte yürürlükteki son oran kullanılır
  Kısmi ödeme        = önce masraflar ve vekalet ücreti, sonra işlemiş faiz, kalan asıl alacaktan düşülür; faiz sonraki günlerde kalan asıl üzerinden işler
  Güncel borç        = kalan asıl alacak + kalan işlemiş faiz + kalan masraf ve vekalet ücreti       (hesap tarihi: --tarih)
  Günlük faiz        = kalan asıl × hesap tarihindeki oran ÷ 100 ÷ gün esası
Oranları siz verirsiniz; örnek dosyadaki oranlar örnektir. Tahsil harcı, cezaevi harcı, icra inkâr tazminatı ve vekalet ücretine işleyen faiz hesaba dahil değildir.
İcra dairesinin kapak hesabı esastır; bu çıktı ön kontrol ve müvekkil bilgilendirmesi içindir.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


DOSYA = {"dosya": ("dosya no", "dosya", "esas no", "takip no"), "borclu": ("borclu", "borclu adi"), "asil": ("asil alacak", "asil", "alacak", "ana para"), "bas": ("faiz baslangic", "faiz baslangici", "faiz baslangic tarihi", "vade", "temerrut tarihi"),
         "tur": ("faiz turu", "faiz", "faiz tipi"), "sabit": ("sabit oran", "akdi oran", "oran"), "masraf": ("masraflar", "masraf", "takip masrafi"), "vekalet": ("vekalet ucreti", "vekalet")}
ODEME = {"dosya": DOSYA["dosya"], "tarih": ("tarih", "odeme tarihi"), "tutar": ("tutar", "odeme", "odenen")}
ORAN = {"tur": DOSYA["tur"], "bas": ("baslangic", "baslangic tarihi", "yururluk", "yururluk tarihi"), "oran": ("yillik oran", "oran", "faiz orani")}
Z, GUN = Decimal(0), timedelta(days=1)


def oku(dosya_yol: Path, oran_yol: Path | None = None, odeme_yol: Path | None = None) -> tuple[list[dict], dict, list[dict], list[dict]]:
    uy, dosyalar, oranlar, odemeler = [], [], {}, []
    for r in kayitlar(dosya_yol, DOSYA, ("dosya", "asil", "bas")):
        d = metin(r["dosya"])
        if not d:
            continue
        asil, bas, sabit, masraf, vekalet = para(r["asil"]), tarih(r["bas"]), (para(r["sabit"]) if metin(r["sabit"]) else None), (para(r["masraf"]) if metin(r["masraf"]) else Z), (para(r["vekalet"]) if metin(r["vekalet"]) else Z)
        if asil is None or asil <= 0 or bas is None or (metin(r["sabit"]) and (sabit is None or sabit < 0)) or masraf is None or vekalet is None or masraf < 0 or vekalet < 0 or (sabit is None and not metin(r["tur"])):
            uy.append(uyari("Yüksek", "Dosya okunamadı", f"{d} (satır {r['_satir']})", f"Asıl alacak '{metin(r['asil'])}', faiz başlangıcı '{metin(r['bas'])}', faiz türü '{metin(r['tur'])}', sabit oran '{metin(r['sabit'])}'. "
                            "Asıl alacak pozitif olmalı; faiz türü veya sabit oran verilmeli. Dosya hesaplanmadı."))
            continue
        dosyalar.append({"dosya": d, "borclu": metin(r["borclu"]), "asil": asil, "bas": bas, "tur": metin(r["tur"]), "sabit": sabit, "masraf": masraf + vekalet})
    if oran_yol:
        for r in kayitlar(oran_yol, ORAN, ("tur", "bas", "oran")):
            t = metin(r["tur"])
            if not t:
                continue
            b, o = tarih(r["bas"]), para(r["oran"])
            if b is None or o is None or o < 0:
                uy.append(uyari("Yüksek", "Oran satırı okunamadı", f"{t} (satır {r['_satir']})", f"Başlangıç '{metin(r['bas'])}', oran '{metin(r['oran'])}'. Satır alınmadı; bu türdeki dosyaların faizi eksik hesaplanabilir."))
                continue
            oranlar.setdefault(katla(t), []).append((b, o))
        for l in oranlar.values():
            l.sort()
    if odeme_yol:
        for r in kayitlar(odeme_yol, ODEME, ("dosya", "tarih", "tutar")):
            d = metin(r["dosya"])
            if not d:
                continue
            t, tu = tarih(r["tarih"]), para(r["tutar"])
            if t is None or tu is None or tu <= 0:
                uy.append(uyari("Yüksek", "Ödeme okunamadı", f"{d} (satır {r['_satir']})", f"Tarih '{metin(r['tarih'])}', tutar '{metin(r['tutar'])}'. Ödeme hesaba katılmadı; borç olduğundan yüksek çıkar."))
                continue
            odemeler.append({"dosya": d, "tarih": t, "tutar": tu})
    return dosyalar, oranlar, odemeler, uy


def oran_bul(tablo: list[tuple], t: date) -> Decimal | None:
    gecerli = [o for b, o in tablo if b <= t]
    return gecerli[-1] if gecerli else None


def hesapla(dosya: dict, oranlar: dict, odemeler: list[dict], hesap: date, gun_esasi: int = 365) -> dict:
    """Tek dosya: dönem dökümü ve güncel borç. Oran bulunamazsa ValueError."""
    tablo = [(date.min, dosya["sabit"])] if dosya["sabit"] is not None else oranlar.get(katla(dosya["tur"]))
    if not tablo:
        raise ValueError(f"'{dosya['tur']}' faiz türü için oran tanımlı değil.")
    if hesap < dosya["bas"]:
        raise ValueError(f"Hesap tarihi ({hesap:%d.%m.%Y}) faiz başlangıcından ({dosya['bas']:%d.%m.%Y}) önce.")
    if oran_bul(tablo, dosya["bas"]) is None:
        raise ValueError(f"'{dosya['tur']}' için {dosya['bas']:%d.%m.%Y} tarihinde yürürlükte oran yok (ilk oran {tablo[0][0]:%d.%m.%Y}).")
    kendi = sorted((o for o in odemeler if katla(o["dosya"]) == katla(dosya["dosya"])), key=lambda o: o["tarih"])
    uy = [uyari("Orta", "Ödeme hesap aralığı dışında", dosya["dosya"], f"{o['tarih']:%d.%m.%Y} tarihli {o['tutar']:,.2f} ödeme faiz başlangıcından önce veya hesap tarihinden sonra; hesaba katılmadı.")
          for o in kendi if not dosya["bas"] <= o["tarih"] <= hesap]
    kendi = [o for o in kendi if dosya["bas"] <= o["tarih"] <= hesap]
    kesim = sorted({b for b, _ in tablo if dosya["bas"] < b < hesap} | {o["tarih"] for o in kendi if dosya["bas"] < o["tarih"] < hesap} | {hesap})
    asil, faiz, masraf, t = dosya["asil"], Z, dosya["masraf"], dosya["bas"]
    donemler, mahsup, top_faiz, top_odeme = [], [], Z, Z

    def ode(gun: date) -> None:
        nonlocal asil, faiz, masraf, top_odeme
        for o in (x for x in kendi if x["tarih"] == gun):
            kalan = o["tutar"]
            m = min(kalan, masraf); kalan -= m; masraf -= m          # noqa: E702
            f = min(kalan, faiz); kalan -= f; faiz -= f              # noqa: E702
            a = min(kalan, asil); kalan -= a; asil -= a              # noqa: E702
            top_odeme += o["tutar"]
            mahsup.append({"dosya": dosya["dosya"], "tarih": gun, "tutar": o["tutar"], "masraf": yuvarla(m), "faiz": yuvarla(f), "asil": yuvarla(a), "fazla": yuvarla(kalan), "kalan_asil": yuvarla(asil)})
            if kalan > 0:
                uy.append(uyari("Orta", "Fazla ödeme", dosya["dosya"], f"{gun:%d.%m.%Y} tarihli ödeme borcun tamamını kapattıktan sonra {kalan:,.2f} artıyor. Tutarları ve dosya eşleşmesini kontrol edin."))

    ode(dosya["bas"])
    for k in kesim:
        o, gun = oran_bul(tablo, t), (k - t).days
        f = asil * o / 100 * gun / gun_esasi
        if gun and asil:
            donemler.append({"dosya": dosya["dosya"], "bas": t, "bit": k, "gun": gun, "asil": yuvarla(asil), "oran": o, "faiz": yuvarla(f)})
        faiz += f
        top_faiz += f
        t = k
        ode(k)
    son_oran = oran_bul(tablo, hesap)
    return {"dosya": dosya["dosya"], "borclu": dosya["borclu"], "tur": dosya["tur"] or f"Sabit %{dosya['sabit'].normalize():f}", "asil": dosya["asil"], "bas": dosya["bas"], "gun": (hesap - dosya["bas"]).days,
            "islemis_faiz": yuvarla(top_faiz), "masraf": dosya["masraf"], "odenen": top_odeme, "kalan_asil": yuvarla(asil), "kalan_faiz": yuvarla(faiz), "kalan_masraf": yuvarla(masraf), "borc": yuvarla(asil + faiz + masraf),
            "oran": son_oran, "gunluk": yuvarla(asil * son_oran / 100 / gun_esasi), "donemler": donemler, "mahsup": mahsup, "uyarilar": uy}


def rapor_yaz(cikti: Path, sonuc: list[dict], hesap: date, gun_esasi: int, uy: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Dosya Özeti"
    _baslik(ws, ["Dosya No", "Borçlu", "Faiz Türü", "Asıl Alacak", "Faiz Başlangıcı", "Gün", "İşlemiş Faiz", "Masraf + Vekalet", "Ödenen", "Kalan Asıl", "Kalan Faiz", "Kalan Masraf", "Güncel Borç", "Güncel Oran %", "Günlük Faiz", "Not (siz doldurun)"],
            (16, 28, 16, 15, 13, 7, 14, 14, 14, 15, 14, 13, 16, 9, 11, 28))
    for s in sonuc:
        _satir(ws, [s["dosya"], s["borclu"], s["tur"], s["asil"], s["bas"], s["gun"], s["islemis_faiz"], s["masraf"] or None, s["odenen"] or None, s["kalan_asil"], s["kalan_faiz"], s["kalan_masraf"] or None, s["borc"], s["oran"], s["gunluk"], ""],
               para_s=(4, 7, 8, 9, 10, 11, 12, 13, 14, 15), tarih_s=(5,), dolgu={13: SARI})
    if sonuc:
        n = _satir(ws, ["Toplam", "", "", sum(s["asil"] for s in sonuc), None, None, sum(s["islemis_faiz"] for s in sonuc), None, sum(s["odenen"] for s in sonuc) or None, sum(s["kalan_asil"] for s in sonuc), sum(s["kalan_faiz"] for s in sonuc), None,
                        sum(s["borc"] for s in sonuc)], para_s=(4, 7, 9, 10, 11, 13))
        for c in ws[n]:
            c.font = Font(bold=True)
    ws.append([])
    ws.append([f"Hesap tarihi {hesap:%d.%m.%Y} · basit faiz, gün esası {gun_esasi} · ödemeler önce masraf ve vekalet ücretine, sonra faize, sonra asıl alacağa mahsup edildi. Harçlar ve tazminatlar dahil değildir; icra dairesinin kapak hesabı esastır."])
    ws = wb.create_sheet("Faiz Dönemleri")
    _baslik(ws, ["Dosya No", "Dönem Başı", "Dönem Sonu", "Gün", "Asıl Alacak", "Yıllık Oran %", "Dönem Faizi"], (16, 13, 13, 7, 16, 11, 15))
    for s in sonuc:
        for d in s["donemler"]:
            _satir(ws, [d["dosya"], d["bas"], d["bit"], d["gun"], d["asil"], d["oran"], d["faiz"]], para_s=(5, 6, 7), tarih_s=(2, 3))
    ws = wb.create_sheet("Ödeme Mahsubu")
    _baslik(ws, ["Dosya No", "Ödeme Tarihi", "Ödeme", "Masrafa", "Faize", "Asıl Alacağa", "Artan", "Kalan Asıl"], (16, 13, 15, 14, 14, 15, 12, 16))
    for s in sonuc:
        for m in s["mahsup"]:
            _satir(ws, [m["dosya"], m["tarih"], m["tutar"], m["masraf"] or None, m["faiz"] or None, m["asil"] or None, m["fazla"] or None, m["kalan_asil"]], para_s=(3, 4, 5, 6, 7, 8), tarih_s=(2,))
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(dosya_yol: Path, cikti: Path, hesap: date, oran_yol: Path | None = None, odeme_yol: Path | None = None, gun_esasi: int = 365) -> dict:
    dosyalar, oranlar, odemeler, uy = oku(dosya_yol, oran_yol, odeme_yol)
    if not dosyalar:
        raise ValueError("Hesaplanacak dosya yok.")
    sonuc, bilinen = [], {katla(d["dosya"]) for d in dosyalar}
    for o in odemeler:
        if katla(o["dosya"]) not in bilinen:
            uy.append(uyari("Orta", "Dosyası bulunamayan ödeme", o["dosya"], f"{o['tarih']:%d.%m.%Y} tarihli {o['tutar']:,.2f} ödemenin dosyası listede yok."))
    for d in dosyalar:
        try:
            s = hesapla(d, oranlar, odemeler, hesap, gun_esasi)
        except ValueError as h:
            uy.append(uyari("Yüksek", "Dosya hesaplanamadı", d["dosya"], str(h)))
            continue
        uy += s.pop("uyarilar")
        sonuc.append(s)
    if not sonuc:
        raise ValueError("Hiçbir dosya hesaplanamadı: " + "; ".join(u["aciklama"] for u in uy if u["tur"] == "Dosya hesaplanamadı")[:300])
    rapor_yaz(cikti, sonuc, hesap, gun_esasi, uy)
    return {"dosyalar": sonuc, "hesap": hesap, "uyarilar": sirala(uy)}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="İcra dosyalarında dönemsel oranlarla işlemiş faizi ve güncel borcu hesaplar.")
    p.add_argument("--dosyalar", type=Path, default=ORNEK / "dosyalar.csv", help="Dosya No, Borçlu, Asıl Alacak, Faiz Başlangıç, Faiz Türü, Sabit Oran, Masraflar, Vekalet Ücreti")
    p.add_argument("--oranlar", type=Path, help="Faiz Türü, Başlangıç, Yıllık Oran")
    p.add_argument("--odemeler", type=Path, help="İsteğe bağlı: Dosya No, Tarih, Tutar")
    p.add_argument("--tarih", help="Hesap tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--gun-esasi", type=int, default=365, choices=(360, 365), help="Yıl gün sayısı (varsayılan 365)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "icra_faiz_hesabi.xlsx")
    a = p.parse_args(argv)
    ornek = a.dosyalar == ORNEK / "dosyalar.csv"
    oran, odeme = a.oranlar or (ORNEK / "faiz_oranlari.csv" if ornek else None), a.odemeler or (ORNEK / "odemeler.csv" if ornek else None)
    hesap = tarih(a.tarih) if a.tarih else date(2026, 10, 13) if ornek else date.today()
    if hesap is None:
        print("[X] --tarih GG.AA.YYYY biçiminde olmalı")
        return 2
    for y in (a.dosyalar, oran, odeme):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.dosyalar, a.cikti, hesap, oran, odeme, a.gun_esasi)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    for s in r["dosyalar"]:
        print(f"[OK] {s['dosya']} {s['borclu']}: asıl {s['kalan_asil']:,.2f} + faiz {s['kalan_faiz']:,.2f} + masraf {s['kalan_masraf']:,.2f} = {s['borc']:,.2f} · günlük faiz {s['gunluk']:,.2f}")
    for u in r["uyarilar"]:
        if u["onem"] != "Bilgi":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Hesap tarihi {hesap:%d.%m.%Y} · Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
