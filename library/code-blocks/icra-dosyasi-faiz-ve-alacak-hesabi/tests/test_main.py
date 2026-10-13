import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
ORAN = {"yasal": [(date(2020, 1, 1), D(24)), (date(2026, 4, 1), D(36))]}
H = date(2026, 7, 1)


def ds(asil=100000, bas=date(2026, 1, 1), tur="Yasal", sabit=None, masraf=0):
    return {"dosya": "D1", "borclu": "B", "asil": D(asil), "bas": bas, "tur": tur, "sabit": None if sabit is None else D(sabit), "masraf": D(masraf)}


def od(ay, gun, tutar, dosya="D1"):
    return {"dosya": dosya, "tarih": date(2026, ay, gun), "tutar": D(tutar)}


class HesapTesti(unittest.TestCase):
    def test_oran_degisimi(self):
        s = main.hesapla(ds(), ORAN, [], H)
        self.assertEqual([(d["bas"], d["bit"], d["gun"], d["oran"], d["faiz"]) for d in s["donemler"]], [(date(2026, 1, 1), date(2026, 4, 1), 90, D(24), D("5917.81")), (date(2026, 4, 1), H, 91, D(36), D("8975.34"))])
        self.assertEqual((s["gun"], s["islemis_faiz"], s["kalan_asil"], s["kalan_faiz"], s["borc"], s["oran"], s["gunluk"], s["uyarilar"]), (181, D("14893.15"), D("100000.00"), D("14893.15"), D("114893.15"), D(36), D("98.63"), []))
        s = main.hesapla(ds(), ORAN, [], H, 360)
        self.assertEqual((s["donemler"][0]["faiz"], s["gunluk"]), (D("6000.00"), D("100.00")))
        s = main.hesapla(ds(sabit=10, tur=""), {}, [], date(2027, 1, 1))
        self.assertEqual((s["tur"], s["gun"], s["islemis_faiz"], len(s["donemler"])), ("Sabit %10", 365, D("10000.00"), 1))
        self.assertEqual(main.hesapla(ds(), ORAN, [], date(2026, 1, 1))["borc"], D("100000.00"))

    def test_kismi_odeme_mahsubu(self):
        s = main.hesapla(ds(masraf=3000), ORAN, [od(3, 1, 10000)], H)
        self.assertEqual(s["mahsup"], [{"dosya": "D1", "tarih": date(2026, 3, 1), "tutar": D(10000), "masraf": D("3000.00"), "faiz": D("3879.45"), "asil": D("3120.55"), "fazla": D("0.00"), "kalan_asil": D("96879.45")}])
        self.assertEqual([(d["gun"], d["asil"], d["faiz"]) for d in s["donemler"]], [(59, D("100000.00"), D("3879.45")), (31, D("96879.45"), D("1974.75")), (91, D("96879.45"), D("8695.26"))])
        self.assertEqual((s["islemis_faiz"], s["odenen"], s["kalan_asil"], s["kalan_faiz"], s["kalan_masraf"], s["borc"], s["gunluk"]), (D("14549.46"), D(10000), D("96879.45"), D("10670.01"), D("0.00"), D("107549.46"), D("95.55")))
        s = main.hesapla(ds(masraf=3000), ORAN, [od(3, 1, 2000), od(3, 1, 500)], H)
        self.assertEqual(([(m["masraf"], m["faiz"], m["asil"]) for m in s["mahsup"]], s["kalan_masraf"], s["kalan_asil"]), ([(D("2000.00"), D("0.00"), D("0.00")), (D("500.00"), D("0.00"), D("0.00"))], D("500.00"), D("100000.00")))
        s = main.hesapla(ds(asil=1000), ORAN, [od(1, 1, 400), od(7, 1, 5000), od(8, 1, 10), od(1, 1, 5, "D2")], H)
        self.assertEqual((s["mahsup"][0]["asil"], s["mahsup"][1]["fazla"] > 0, s["borc"], sorted(u["tur"] for u in s["uyarilar"])), (D("400.00"), True, D("0.00"), ["Fazla ödeme", "Ödeme hesap aralığı dışında"]))

    def test_hatalar(self):
        for dosya, hesap in ((ds(tur="Bilinmeyen"), H), (ds(bas=date(2019, 6, 1)), H), (ds(), date(2025, 12, 31))):
            with self.assertRaises(ValueError):
                main.hesapla(dosya, ORAN, [], hesap)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "i.xlsx"
            r = main.calistir(ORNEK / "dosyalar.csv", cikti, date(2026, 10, 13), ORNEK / "faiz_oranlari.csv", ORNEK / "odemeler.csv")
            self.assertEqual([(s["dosya"], s["gun"], s["islemis_faiz"], s["kalan_asil"], s["borc"], s["gunluk"]) for s in r["dosyalar"]][1:],
                             [("2026/102", 212, D("55178.08"), D("250000.00"), D("305178.08"), D("260.27")), ("2026/103", 134, D("5506.85"), D("50000.00"), D("55506.85"), D("41.10"))])
            a = r["dosyalar"][0]
            self.assertEqual((a["kalan_asil"], a["kalan_masraf"], a["odenen"], a["oran"], r["uyarilar"]), (D("96879.45"), D("0.00"), D("10000.00"), D(24), []))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Dosya Özeti", "Faiz Dönemleri", "Ödeme Mahsubu", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            dy, oy, py = Path(tmp) / "d.csv", Path(tmp) / "o.csv", Path(tmp) / "p.csv"
            dy.write_text("Dosya;Borçlu;Alacak;Vade;Faiz;Oran;Masraf;Vekalet\nA;X;1.250,50;01.01.2026;Yasal;;100;50\nB;Y;0;01.01.2026;Yasal;;;\nC;Z;100;01.01.2026;;;;\nD;W;100;01.01.2026;Akdi;;;\nE;V;100;01.01.2026;;12,5;;\n", encoding="utf-8")
            oy.write_text("Faiz Türü;Başlangıç;Oran\nYasal;01.06.2024;24\nYasal;01.01.2020;%9\nYasal;x;5\n", encoding="utf-8")
            py.write_text("Dosya;Tarih;Tutar\nA;01.02.2026;200\nA;x;5\nQ;01.02.2026;10\n", encoding="utf-8")
            d, o, p, uy = main.oku(dy, oy, py)
            self.assertEqual(([(x["dosya"], x["asil"], x["sabit"], x["masraf"]) for x in d], o, len(p), sorted(u["tur"] for u in uy)),
                             ([("A", D("1250.50"), None, D(150)), ("D", D(100), None, D(0)), ("E", D(100), D("12.5"), D(0))], {"yasal": [(date(2020, 1, 1), D(9)), (date(2024, 6, 1), D(24))]}, 2,
                              ["Dosya okunamadı", "Dosya okunamadı", "Oran satırı okunamadı", "Ödeme okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            r = main.calistir(dy, Path(c), date(2026, 3, 1), oy, py)
            self.assertEqual(([s["dosya"] for s in r["dosyalar"]], sorted(u["tur"] for u in r["uyarilar"])), (["A", "E"], ["Dosya hesaplanamadı", "Dosya okunamadı", "Dosya okunamadı", "Dosyası bulunamayan ödeme", "Oran satırı okunamadı", "Ödeme okunamadı"]))
            arg = ["--dosyalar", str(dy), "--oranlar", str(oy), "--cikti", c]
            self.assertEqual(main.main(arg + ["--odemeler", str(py), "--tarih", "01.03.2026", "--gun-esasi", "360"]), 0)
            self.assertEqual(main.main(arg + ["--tarih", "x"]), 2)
            self.assertEqual(main.main(arg + ["--tarih", "01.01.2025"]), 1)
            self.assertEqual(main.main(arg + ["--odemeler", str(Path(tmp) / "yok.csv")]), 1)
            dy.write_text("Dosya;Borçlu\nA;X\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
