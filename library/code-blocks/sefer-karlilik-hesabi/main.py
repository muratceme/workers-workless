"""
Sefer Kârlılık Hesabı — Gerisi Bende kod bloğu
Karayolu Operasyon › Karayolu Operasyon Uzmanı

Navlun geliri, alt nakliyeci bedeli, otoyol, gümrük ve diğer masraflardan sefer bazında kârlılığı hesaplar.
  Toplam masraf   = alt nakliyeci + yakıt + otoyol / geçiş + gümrük + harcırah + diğer
  Kâr             = navlun geliri − toplam masraf;        marj % = kâr ÷ navlun geliri × 100
  Km başına       = gelir ÷ km, masraf ÷ km, kâr ÷ km     (km yazılmışsa)
  İşaret          = Zarar (kâr < 0) · Düşük marj (0 ≤ marj < --marj-esik, varsayılan %10)
  Özetler         = müşteri ve güzergah bazında toplam gelir, masraf, kâr, marj
Sefere doğrudan yazılan masraflar hesaplanır; amortisman, sigorta, bakım, sürücü maaşı ve genel giderler girilmediyse kâr olduğundan yüksek görünür.
Tüm tutarlar aynı para biriminde ve aynı KDV esasında olmalıdır; kur çevirisi yapılmaz.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


MASRAF = (("alt", "Alt Nakliyeci"), ("yakit", "Yakıt"), ("otoyol", "Otoyol / Geçiş"), ("gumruk", "Gümrük"), ("harcirah", "Harcırah"), ("diger", "Diğer"))
SUTUN = {"no": ("sefer no", "sefer", "sefer numarasi", "dosya no"), "tarih": ("tarih", "sefer tarihi", "yukleme tarihi"), "musteri": ("musteri", "gonderici", "cari"), "guzergah": ("guzergah", "hat", "rota"),
         "km": ("km", "mesafe", "toplam km"), "gelir": ("navlun geliri", "navlun", "gelir", "satis"), "alt": ("alt nakliyeci", "alt nakliyeci bedeli", "tasimaci bedeli", "nakliyeci"),
         "yakit": ("yakit", "yakit gideri", "akaryakit"), "otoyol": ("otoyol", "otoyol gecis", "otoyol kopru", "gecis ucreti", "yol gecis"), "gumruk": ("gumruk", "gumruk masrafi"),
         "harcirah": ("harcirah", "surucu harcirahi", "surucu"), "diger": ("diger", "diger masraf", "diger masraflar")}
Z = Decimal(0)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("no", "gelir")):
        no = metin(r["no"])
        if not no:
            continue
        kim = f"{no} (satır {r['_satir']})"
        gelir = para(r["gelir"])
        if gelir is None or gelir < 0:
            uy.append(uyari("Yüksek", "Navlun geliri okunamadı", kim, f"'{metin(r['gelir'])}'. Sefer alınmadı."))
            continue
        x = {"no": no, "tarih": tarih(r["tarih"]), "musteri": metin(r["musteri"]) or "Belirtilmemiş", "guzergah": metin(r["guzergah"]) or "Belirtilmemiş", "gelir": gelir}
        for a, ad in (*MASRAF, ("km", "Km")):
            v = para(r[a])
            if v is None and metin(r[a]) or (v is not None and v < 0):
                uy.append(uyari("Yüksek" if a != "km" else "Orta", "Değer okunamadı", kim, f"{ad} '{metin(r[a])}'. 0 sayıldı; sefer kârı olduğundan yüksek görünebilir."))
                v = None
            x[a] = v or Z
        if gelir == 0:
            uy.append(uyari("Orta", "Geliri olmayan sefer", kim, "Navlun geliri 0: boş dönüş veya henüz faturalanmamış sefer olabilir. Marj hesaplanmadı."))
        if sum((x[a] for a, _ in MASRAF), Z) == 0:
            uy.append(uyari("Orta", "Masrafı olmayan sefer", kim, "Hiç masraf yazılmamış; kâr gelire eşit görünür."))
        l.append(x)
    return l, uy


def olcut(x: dict) -> dict:
    masraf = sum((x[a] for a, _ in MASRAF), Z)
    kar, km = x["gelir"] - masraf, x["km"]
    return {"masraf": masraf, "kar": kar, "marj": yuvarla(kar / x["gelir"] * 100) if x["gelir"] else None, "km_gelir": yuvarla(x["gelir"] / km) if km else None,
            "km_masraf": yuvarla(masraf / km) if km else None, "km_kar": yuvarla(kar / km) if km else None}


def hesapla(seferler: list[dict], marj_esik: Decimal = Decimal(10)) -> dict:
    if not seferler:
        raise ValueError("Hesaplanacak sefer yok.")
    for s in seferler:
        s |= olcut(s)
        s["isaret"] = "Zarar" if s["kar"] < 0 else "Düşük marj" if s["marj"] is not None and s["marj"] < marj_esik else ""

    def grupla(alan) -> list[dict]:
        g = {}
        for s in seferler:
            x = g.setdefault(katla(s[alan]) if alan else "", {"ad": s[alan] if alan else "Toplam", "sefer": 0, "zarar": 0, "gelir": Z, "km": Z} | {a: Z for a, _ in MASRAF})
            x["sefer"] += 1
            x["zarar"] += s["kar"] < 0
            for a in ("gelir", "km", *[m for m, _ in MASRAF]):
                x[a] += s[a]
        return sorted((x | olcut(x) for x in g.values()), key=lambda x: (x["kar"], x["ad"]))

    return {"seferler": sorted(seferler, key=lambda s: (s["kar"], s["no"])), "musteri": grupla("musteri"), "guzergah": grupla("guzergah"), "genel": grupla(None)[0], "esik": marj_esik,
            "say": {"Zarar": sum(s["isaret"] == "Zarar" for s in seferler), "Düşük marj": sum(s["isaret"] == "Düşük marj" for s in seferler)}}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    mb = [ad for _, ad in MASRAF]
    son = ["Toplam Masraf", "Kâr", "Marj %", "Km", "Gelir / Km", "Masraf / Km", "Kâr / Km"]
    deger = lambda x: [x["gelir"], *[x[a] for a, _ in MASRAF], x["masraf"], x["kar"], x["marj"], x["km"] or None, x["km_gelir"], x["km_masraf"], x["km_kar"]]   # noqa: E731
    ws = wb.active
    ws.title = "Sefer Bazında"
    _baslik(ws, ["İşaret", "Sefer No", "Tarih", "Müşteri", "Güzergah", "Navlun Geliri", *mb, *son], (12, 14, 12, 24, 24, 14, 13, 12, 12, 12, 12, 12, 14, 14, 9, 9, 10, 10, 10))
    for s in r["seferler"]:
        _satir(ws, [s["isaret"], s["no"], s["tarih"], s["musteri"], s["guzergah"], *deger(s)], para_s=tuple(range(6, 20)), tarih_s=(3,),
               dolgu={1: KIRMIZI if s["isaret"] == "Zarar" else SARI if s["isaret"] else None, 14: KIRMIZI if s["kar"] < 0 else None})
    g = r["genel"]
    _satir(ws, ["", "Toplam", None, "", "", *deger(g)], para_s=tuple(range(6, 20)))
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.append([])
    ws.append([f"Düşük marj eşiği %{r['esik'].normalize():f}. Yalnız sefere yazılan masraflar hesaba girer; amortisman, sigorta, bakım ve genel giderler girilmediyse kâr olduğundan yüksek görünür."])
    for ad, anahtar, bas in (("Müşteri Bazında", "musteri", "Müşteri"), ("Güzergah Bazında", "guzergah", "Güzergah")):
        ws = wb.create_sheet(ad)
        _baslik(ws, [bas, "Sefer", "Zarar Eden Sefer", "Navlun Geliri", *mb, *son], (26, 8, 9, 14, 13, 12, 12, 12, 12, 12, 14, 14, 9, 9, 10, 10, 10))
        for x in r[anahtar]:
            _satir(ws, [x["ad"], x["sefer"], x["zarar"] or None, *deger(x)], para_s=tuple(range(4, 18)), dolgu={12: KIRMIZI if x["kar"] < 0 else None})
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, marj_esik: Decimal = Decimal(10)) -> dict:
    seferler, uy = oku(yol)
    r = hesapla(seferler, marj_esik)
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Sefer bazında navlun geliri ve masraflardan kârlılık hesaplar.")
    p.add_argument("--seferler", type=Path, default=ORNEK / "seferler.csv", help="Sefer No, Tarih, Müşteri, Güzergah, Km, Navlun Geliri, Alt Nakliyeci, Yakıt, Otoyol, Gümrük, Harcırah, Diğer")
    p.add_argument("--marj-esik", default="10", help="Bu marjın (%%) altındaki kârlı seferler 'Düşük marj' (varsayılan 10)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "sefer_karlilik.xlsx")
    a = p.parse_args(argv)
    e = para(a.marj_esik)
    if e is None or e < 0:
        print("[X] --marj-esik sıfır veya pozitif sayı olmalı")
        return 2
    if not a.seferler.exists():
        print(f"[X] Dosya bulunamadı: {a.seferler}")
        return 1
    try:
        r = calistir(a.seferler, a.cikti, e)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    g, s = r["genel"], r["say"]
    print(f"[OK] {g['sefer']} sefer · gelir {g['gelir']:,.2f} · masraf {g['masraf']:,.2f} · kâr {g['kar']:,.2f}" + (f" (marj %{g['marj']})" if g["marj"] is not None else "") + f" · {s['Zarar']} zarar, {s['Düşük marj']} düşük marj")
    for x in r["seferler"]:
        if x["isaret"]:
            print(f"[!] {x['no']} · {x['musteri']} · {x['guzergah']}: {x['isaret']}, kâr {x['kar']:,.2f}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
