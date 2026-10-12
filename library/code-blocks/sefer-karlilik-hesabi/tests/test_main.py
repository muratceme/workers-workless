import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def sf(no, gelir, km=0, musteri="M", guzergah="G", **masraf):
    return {"no": no, "tarih": None, "musteri": musteri, "guzergah": guzergah, "gelir": D(gelir), "km": D(km)} | {a: D(masraf.get(a, 0)) for a, _ in main.MASRAF}


def tek(*a, **kw):
    esik = kw.pop("esik", D(10))
    return main.hesapla([sf("S", *a, **kw)], esik)["seferler"][0]


class HesapTesti(unittest.TestCase):
    def test_sefer(self):
        s = tek(3000, 2000, yakit=1100, otoyol=300, gumruk=150, harcirah=250, diger=100)
        self.assertEqual((s["masraf"], s["kar"], s["marj"], s["km_gelir"], s["km_masraf"], s["km_kar"], s["isaret"]), (D(1900), D(1100), D("36.67"), D("1.50"), D("0.95"), D("0.55"), ""))
        s = tek(0, alt=500)
        self.assertEqual((s["kar"], s["marj"], s["km_gelir"], s["isaret"]), (D(-500), None, None, "Zarar"))
        self.assertEqual([tek(1000, alt=m)["isaret"] for m in (900, 901, 1000, "1000.01")], ["", "Düşük marj", "Düşük marj", "Zarar"])        # tam %10 işaretlenmez; kâr 0 düşük marj
        self.assertEqual(tek(1000, alt=850, esik=D(20))["isaret"], "Düşük marj")
        with self.assertRaises(ValueError):
            main.hesapla([])

    def test_ozetler(self):
        l = [sf("1", 3000, 2000, "A", "X", yakit=1900), sf("2", 2800, 2000, "A", "X", alt=2750), sf("3", 3200, 2300, "B", "Y", diger=3300), sf("4", 3400, 2300, "b", "Y", yakit=2200)]
        r = main.hesapla(l)
        self.assertEqual([s["no"] for s in r["seferler"]], ["3", "2", "1", "4"])
        self.assertEqual([(x["ad"], x["sefer"], x["zarar"], x["gelir"], x["kar"], x["marj"]) for x in r["musteri"]], [("B", 2, 1, D(6600), D(1100), D("16.67")), ("A", 2, 0, D(5800), D(1150), D("19.83"))])
        g = r["genel"]
        self.assertEqual((g["gelir"], g["masraf"], g["kar"], g["marj"], g["km"], g["km_gelir"], g["km_masraf"], r["say"]), (D(12400), D(10150), D(2250), D("18.15"), D(8600), D("1.44"), D("1.18"), {"Zarar": 1, "Düşük marj": 1}))
        self.assertEqual([(x["ad"], x["yakit"], x["alt"]) for x in r["guzergah"]], [("Y", D(2200), D(0)), ("X", D(1900), D(2750))])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "s.xlsx"
            r = main.calistir(ORNEK / "seferler.csv", cikti)
            self.assertEqual([(s["no"], s["masraf"], s["kar"], s["marj"], s["isaret"]) for s in r["seferler"]],
                             [("S-03", D(3300), D(-100), D("-3.13"), "Zarar"), ("S-02", D(2750), D(50), D("1.79"), "Düşük marj"), ("S-01", D(1900), D(1100), D("36.67"), ""), ("S-04", D(2200), D(1200), D("35.29"), "")])
            self.assertEqual((r["genel"]["kar"], r["genel"]["marj"], r["uyarilar"], load_workbook(cikti).sheetnames), (D(2250), D("18.15"), [], ["Sefer Bazında", "Müşteri Bazında", "Güzergah Bazında", "Uyarılar"]))

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "s.csv"
            y.write_text("Sefer No;Navlun;Yakıt;Km\nA;1.250,50;çok;2.000\nB;yok;1;1\nC;0;;\nD;500;100;uzak\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual([(s["no"], s["gelir"], s["yakit"], s["km"], s["musteri"]) for s in l], [("A", D("1250.50"), D(0), D(2000), "Belirtilmemiş"), ("C", D(0), D(0), D(0), "Belirtilmemiş"), ("D", D(500), D(100), D(0), "Belirtilmemiş")])
            self.assertEqual(sorted(u["tur"] for u in uy), ["Değer okunamadı", "Değer okunamadı", "Geliri olmayan sefer", "Masrafı olmayan sefer", "Masrafı olmayan sefer", "Navlun geliri okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--seferler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--marj-esik", "15"]), 0)
            self.assertEqual(main.main(arg + ["--marj-esik", "-1"]), 2)
            self.assertEqual(main.main(["--seferler", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Sefer No;Navlun Geliri\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
