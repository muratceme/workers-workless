"""
Kullanıcı Yetki Gözden Geçirme — Workers / Workless kod bloğu
Bilgi Teknolojileri › Bilgi Güvenliği Uzmanı

Sistemlerdeki kullanıcı-yetki listelerini İK personel listesiyle karşılaştırır (periyodik erişim gözden geçirmesi):
  - Eşleştirme sırası: sicil no → e-posta → ad soyad (Türkçe karakter ve büyük/küçük harf duyarsız).
  - Ayrılmış personelin aktif hesabı; ayrılış tarihinden sonra yapılmış giriş (olay incelemesi gerektirir);
    kilitli / pasif ama silinmemiş hesap.
  - İK listesinde karşılığı olmayan kişisel hesap (sahipsiz); ortak ve servis hesapları (sahip / gerekçe istenir).
  - Uzun süredir kullanılmayan aktif hesap (--pasif-gun, varsayılan 90); izindeki personelin hesabı.
  - Rol matrisi (pozisyon × sistem → izinli roller, "*" tüm pozisyonlar) dışındaki yetkiler; kritik roller.
  - Görevler ayrılığı (SoD): aynı kişide çakışan iki rol (sistemler arası da); rol adında "*" joker kullanılabilir.
Rapor: bulgular, yönetici onay listesi (Onay / İptal sütunuyla), görevler ayrılığı, kullanıcı-hesap eşleşmesi,
sistem özeti. İnternete bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: 11 personel, 3 sistem, 25 yetki satırı
    python main.py --personel ik.xlsx --yetkiler yetkiler.xlsx --rol-matrisi matris.csv --gorev-ayriligi sod.csv
    python main.py --personel ik.xlsx --yetkiler y.xlsx --bugun 31.12.2026 --pasif-gun 60
"""
from __future__ import annotations

import argparse
import csv
import fnmatch
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")

PERSONEL_SUTUNLARI = {"sicil": ("sicil no", "sicil"), "ad": ("ad soyad", "ad"), "eposta": ("e posta", "eposta", "email", "mail"),
                      "departman": ("departman", "birim"), "pozisyon": ("pozisyon", "unvan"), "durum": ("durum", "calisma durumu"),
                      "ayrilis": ("ayrilis tarihi", "cikis tarihi"), "yonetici": ("yonetici", "bagli oldugu yonetici")}
YETKI_SUTUNLARI = {"sistem": ("sistem", "uygulama"), "kullanici": ("kullanici adi", "kullanici", "hesap"), "ad": ("ad soyad", "ad"),
                   "eposta": ("e posta", "eposta", "email", "mail"), "sicil": ("sicil no", "sicil"), "rol": ("rol", "yetki", "rol yetki"),
                   "son_giris": ("son giris", "son giris tarihi", "last login"), "durum": ("hesap durumu", "durum"), "tur": ("hesap turu", "tur")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9@.*]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c).replace("@", " ").replace(".", " ").strip() for c in r]
        b = [" ".join(x.split()) for x in b]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Personel:
    sicil: str
    ad: str
    eposta: str
    departman: str
    pozisyon: str
    durum: str
    ayrilis: date | None
    yonetici: str

    @property
    def ayrildi(self) -> bool:
        return katla(self.durum).startswith(("ayril", "cikis", "pasif")) or self.ayrilis is not None


@dataclass
class Hesap:
    sistem: str
    kullanici: str
    ad: str
    eposta: str
    sicil: str
    roller: list
    son_giris: date | None
    durum: str
    tur: str
    kisi: Personel | None = None
    eslesme: str = ""
    bulgular: list = field(default_factory=list)

    @property
    def aktif(self) -> bool:
        return not katla(self.durum).startswith(("kilit", "pasif", "devre", "disab", "kapal"))

    @property
    def tur_k(self) -> str:
        k = katla(self.tur)
        return "servis" if k.startswith(("servis", "sistem", "service")) else "ortak" if k.startswith(("ortak", "paylasim", "shared", "genel")) else "kisisel"


def oku(personel_yolu: Path, yetki_yolu: Path) -> tuple[list[Personel], list[Hesap]]:
    personel = []
    for r in kayitlar(personel_yolu, PERSONEL_SUTUNLARI, ("ad",)):
        if metin(r.get("ad")):
            personel.append(Personel(metin(r.get("sicil")), metin(r["ad"]), metin(r.get("eposta")).lower(), metin(r.get("departman")),
                                     metin(r.get("pozisyon")), metin(r.get("durum")) or "Aktif", tarih(r.get("ayrilis")), metin(r.get("yonetici"))))
    hesaplar: dict[tuple, Hesap] = {}
    for r in kayitlar(yetki_yolu, YETKI_SUTUNLARI, ("sistem", "kullanici")):
        s, k = metin(r.get("sistem")), metin(r.get("kullanici"))
        if not s or not k:
            continue
        h = hesaplar.get((katla(s), k.lower()))
        if h is None:
            h = Hesap(s, k, metin(r.get("ad")), metin(r.get("eposta")).lower(), metin(r.get("sicil")), [], tarih(r.get("son_giris")), metin(r.get("durum")) or "Aktif",
                      metin(r.get("tur")) or "Kişisel")
            hesaplar[(katla(s), k.lower())] = h
        rol = metin(r.get("rol"))
        if rol and rol not in h.roller:
            h.roller.append(rol)
        sg = tarih(r.get("son_giris"))
        if sg and (not h.son_giris or sg > h.son_giris):
            h.son_giris = sg
    return personel, list(hesaplar.values())


def esle(personel: list[Personel], hesaplar: list[Hesap]) -> None:
    sicil = {p.sicil: p for p in personel if p.sicil}
    eposta = {p.eposta: p for p in personel if p.eposta}
    ad = defaultdict(list)
    for p in personel:
        ad[katla(p.ad)].append(p)
    for h in hesaplar:
        if h.sicil and h.sicil in sicil:
            h.kisi, h.eslesme = sicil[h.sicil], "Sicil"
        elif h.eposta and h.eposta in eposta:
            h.kisi, h.eslesme = eposta[h.eposta], "E-posta"
        elif h.ad and len(ad.get(katla(h.ad), [])) == 1:
            h.kisi, h.eslesme = ad[katla(h.ad)][0], "Ad soyad"
        elif h.ad and len(ad.get(katla(h.ad), [])) > 1:
            h.eslesme = "Belirsiz (aynı ad)"


def rol_uyar(desen: str, rol: str) -> bool:
    return fnmatch.fnmatchcase(katla(rol), katla(desen))


# ----------------------------------------------------------------------------
# Analiz
# ----------------------------------------------------------------------------

def analiz_et(personel, hesaplar, matris, sod, kritik, bugun: date, pasif_gun: int) -> dict:
    esle(personel, hesaplar)
    for h in hesaplar:
        b = h.bulgular
        p = h.kisi
        gun = (bugun - h.son_giris).days if h.son_giris else None
        if h.tur_k == "servis":
            b.append(("Bilgi", "Servis hesabı", "Sahibi, kullanım amacı ve parola / anahtar yönetimi belgelenmeli"))
        elif h.tur_k == "ortak":
            b.append(("Orta", "Ortak hesap", "Birden çok kişinin kullandığı hesapta işlem sorumluluğu izlenemez; kişisel hesaplara geçilmeli"))
        elif p is None:
            if h.eslesme.startswith("Belirsiz"):
                b.append(("Orta", "Belirsiz eşleşme", f"'{h.ad}' adında birden çok personel var; sicil veya e-posta ekleyin"))
            elif h.aktif:
                b.append(("Yüksek", "Sahipsiz hesap", "İK listesinde karşılığı yok (eski çalışan, dış kaynak veya tanımsız kullanıcı olabilir)"))
            else:
                b.append(("Bilgi", "Sahipsiz pasif hesap", "İK listesinde karşılığı yok; hesap pasif, silinmesi önerilir"))
        if p and p.ayrildi:
            if h.aktif:
                b.append(("Yüksek", "Ayrılan personelin aktif hesabı", f"{p.ad} {p.ayrilis:%d.%m.%Y} tarihinde ayrıldı; hesap hâlâ aktif" if p.ayrilis
                          else f"{p.ad} ayrılmış görünüyor; hesap hâlâ aktif"))
            else:
                b.append(("Bilgi", "Ayrılan personelin pasif hesabı", "Hesap kilitli / pasif; saklama politikasına göre silinmeli"))
            if p.ayrilis and h.son_giris and h.son_giris > p.ayrilis:
                b.append(("Yüksek", "Ayrılıştan sonra giriş", f"Son giriş {h.son_giris:%d.%m.%Y}, ayrılış {p.ayrilis:%d.%m.%Y}; olay olarak incelenmeli"))
        elif h.aktif and gun is not None and gun > pasif_gun:
            b.append(("Orta", "Kullanılmayan hesap", f"{gun} gündür giriş yok (> {pasif_gun}); ihtiyaç teyit edilmeli veya kapatılmalı"))
        elif h.aktif and h.son_giris is None:
            b.append(("Bilgi", "Giriş bilgisi yok", "Son giriş tarihi verilmemiş"))
        if p and katla(p.durum).startswith("izin"):
            b.append(("Bilgi", "İzindeki personel", "Uzun izin süresince hesabın askıya alınması değerlendirilebilir"))
        # Rol matrisi
        if p and matris and h.aktif and not p.ayrildi:
            izinli = [r for (poz, sis), rs in matris.items() if katla(sis) == katla(h.sistem) and poz in ("*", katla(p.pozisyon)) for r in rs]
            disinda = [r for r in h.roller if not any(rol_uyar(i, r) for i in izinli)]
            if disinda:
                b.append(("Orta", "Matris dışı yetki", f"{p.pozisyon} için {h.sistem} rol matrisinde yok: {', '.join(disinda)}"))
        h.kritik = [r for r in h.roller if any(rol_uyar(k, r) for k in kritik)]
    # Görevler ayrılığı — kişi bazında, sistemler arası
    kisi_rol = defaultdict(list)
    for h in hesaplar:
        if h.kisi and h.aktif:
            for r in h.roller:
                kisi_rol[h.kisi.sicil or h.kisi.ad].append((h.sistem, r, h))
    ihlaller = []
    for anahtar, lst in kisi_rol.items():
        for ad, a, bdesen, risk in sod:
            ra = [(s, r, h) for s, r, h in lst if rol_uyar(a, r)]
            rb = [(s, r, h) for s, r, h in lst if rol_uyar(bdesen, r)]
            if ra and rb and any(x[2:] != y[2:] or x[1] != y[1] for x in ra for y in rb):
                p = lst[0][2].kisi
                ihlaller.append({"kisi": p, "kural": ad, "a": ", ".join(sorted({f"{s}: {r}" for s, r, _ in ra})),
                                 "b": ", ".join(sorted({f"{s}: {r}" for s, r, _ in rb})), "risk": risk})
                for _, _, h in ra[:1]:
                    h.bulgular.append(("Yüksek", "Görevler ayrılığı", f"{ad}: {ihlaller[-1]['a']} + {ihlaller[-1]['b']}"))
    bulgular = []
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    for h in hesaplar:
        for o, t, a in h.bulgular:
            bulgular.append({"onem": o, "tur": t, "sistem": h.sistem, "kullanici": h.kullanici, "kisi": h.kisi.ad if h.kisi else h.ad,
                             "roller": ", ".join(h.roller), "aciklama": a})
    for p in personel:
        if p.ayrildi:
            continue
        if not any(h.kisi is p for h in hesaplar):
            bulgular.append({"onem": "Bilgi", "tur": "Hesabı yok", "sistem": "", "kullanici": "", "kisi": p.ad, "roller": "",
                             "aciklama": "Aktif personelin hiçbir sistemde hesabı listelenmedi (liste eksik olabilir)"})
    bulgular.sort(key=lambda b: (sira[b["onem"]], b["tur"], b["sistem"], b["kullanici"]))
    return {"personel": personel, "hesaplar": hesaplar, "ihlaller": ihlaller, "bulgular": bulgular, "bugun": bugun}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KONTROL = PatternFill("solid", fgColor="FFF4CE")
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, basliklar, genislik=()):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def rapor_yaz(cikti: Path, s: dict) -> None:
    wb = Workbook()
    bl = wb.active
    bl.title = "Bulgular"
    _baslik(bl, ["Önem", "Tür", "Sistem", "Kullanıcı", "Kişi", "Roller", "Açıklama", "Aksiyon", "Tamamlandı"], (9, 28, 14, 14, 18, 34, 60, 20, 11))
    for b in s["bulgular"]:
        bl.append([b["onem"], b["tur"], b["sistem"], b["kullanici"], b["kisi"], b["roller"], b["aciklama"], "", ""])
        bl.cell(bl.max_row, 1).fill = PatternFill("solid", fgColor=RENK[b["onem"]])
        bl.cell(bl.max_row, 7).alignment = UST
        bl.cell(bl.max_row, 8).fill = bl.cell(bl.max_row, 9).fill = KONTROL
    bl.auto_filter.ref = f"A1:I{bl.max_row}"

    yo = wb.create_sheet("Yönetici Onay Listesi")
    _baslik(yo, ["Onaylayacak Yönetici", "Kişi", "Pozisyon", "Sistem", "Kullanıcı", "Rol", "Kritik", "Son Giriş", "Bulgu", "Karar (Onay / İptal)", "Not"],
            (20, 18, 20, 14, 14, 28, 7, 11, 30, 16, 20))
    satirlar = []
    for h in s["hesaplar"]:
        if not h.aktif:
            continue
        yon = h.kisi.yonetici if h.kisi and h.kisi.yonetici else ("Bilgi İşlem (sahipsiz)" if not h.kisi else "Tanımsız")
        for r in h.roller:
            satirlar.append((yon, h.kisi.ad if h.kisi else h.ad or h.kullanici, h.kisi.pozisyon if h.kisi else "", h.sistem, h.kullanici, r,
                             "Evet" if r in h.kritik else "", h.son_giris, "; ".join(sorted({t for o, t, _ in h.bulgular if o != "Bilgi"}))))
    for x in sorted(satirlar, key=lambda x: (x[0], x[1], x[3])):
        yo.append(list(x) + ["", ""])
        r = yo.max_row
        yo.cell(r, 8).number_format = "DD.MM.YYYY"
        if x[6]:
            yo.cell(r, 7).fill = PatternFill("solid", fgColor="FFF4CE")
        if x[8]:
            yo.cell(r, 9).fill = PatternFill("solid", fgColor="FDE2E1")
        yo.cell(r, 10).fill = KONTROL
    yo.auto_filter.ref = f"A1:K{yo.max_row}"

    sd = wb.create_sheet("Görevler Ayrılığı")
    _baslik(sd, ["Kişi", "Pozisyon", "Kural", "Rol A", "Rol B", "Risk", "Karar / Telafi Edici Kontrol"], (18, 20, 30, 34, 34, 50, 28))
    for x in s["ihlaller"]:
        sd.append([x["kisi"].ad, x["kisi"].pozisyon, x["kural"], x["a"], x["b"], x["risk"], ""])
        sd.cell(sd.max_row, 7).fill = KONTROL
        for j in (4, 5, 6):
            sd.cell(sd.max_row, j).alignment = UST

    ek = wb.create_sheet("Hesap Eşleşmesi")
    _baslik(ek, ["Sistem", "Kullanıcı", "Hesap Adı", "Hesap Türü", "Hesap Durumu", "Roller", "Son Giriş", "Eşleşme", "Sicil", "Personel", "Departman",
                 "Personel Durumu"], (14, 14, 18, 10, 11, 34, 11, 12, 8, 18, 14, 12))
    for h in sorted(s["hesaplar"], key=lambda h: (h.sistem, h.kullanici)):
        p = h.kisi
        ek.append([h.sistem, h.kullanici, h.ad, h.tur, h.durum, ", ".join(h.roller), h.son_giris, h.eslesme or "Yok", p.sicil if p else "", p.ad if p else "",
                   p.departman if p else "", p.durum if p else ""])
        ek.cell(ek.max_row, 7).number_format = "DD.MM.YYYY"
        if not p and h.tur_k == "kisisel":
            ek.cell(ek.max_row, 8).fill = PatternFill("solid", fgColor="FDE2E1")
    ek.auto_filter.ref = f"A1:L{ek.max_row}"

    oz = wb.create_sheet("Sistem Özeti")
    _baslik(oz, ["Sistem", "Hesap", "Aktif", "Kritik Rollü", "Yüksek Bulgu", "Orta Bulgu"], (16, 8, 8, 11, 12, 11))
    for sis in sorted({h.sistem for h in s["hesaplar"]}):
        hs = [h for h in s["hesaplar"] if h.sistem == sis]
        oz.append([sis, len(hs), sum(1 for h in hs if h.aktif), sum(1 for h in hs if h.kritik and h.aktif),
                   sum(1 for b in s["bulgular"] if b["sistem"] == sis and b["onem"] == "Yüksek"),
                   sum(1 for b in s["bulgular"] if b["sistem"] == sis and b["onem"] == "Orta")])
    oz.append([])
    oz.append([f"Gözden geçirme tarihi: {s['bugun']:%d.%m.%Y}. Yönetici onay listesi ilgili yöneticilere gönderilip imzalı / onaylı olarak saklanmalıdır."])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(personel_yolu: Path, yetki_yolu: Path, cikti: Path, bugun: date, matris_yolu: Path | None = None, sod_yolu: Path | None = None,
             kritik_yolu: Path | None = None, pasif_gun: int = 90) -> dict:
    personel, hesaplar = oku(personel_yolu, yetki_yolu)
    if not hesaplar:
        raise ValueError(f"{yetki_yolu.name}: yetki kaydı bulunamadı")
    matris = defaultdict(list)
    if matris_yolu:
        for r in tablo_oku(matris_yolu)[1:]:
            if len(r) >= 3 and r[0] and r[1]:
                matris[(katla(r[0]) if metin(r[0]) != "*" else "*", metin(r[1]))] += [x.strip() for x in str(r[2] or "").split("|") if x.strip()]
    sod = [(metin(r[0]), metin(r[1]), metin(r[2]), metin(r[3]) if len(r) > 3 else "") for r in (tablo_oku(sod_yolu)[1:] if sod_yolu else []) if len(r) >= 3]
    kritik = [metin(r[0]) for r in (tablo_oku(kritik_yolu)[1:] if kritik_yolu else []) if r and metin(r[0])]
    s = analiz_et(personel, hesaplar, dict(matris), sod, kritik, bugun, pasif_gun)
    rapor_yaz(cikti, s)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Sistem kullanıcı-yetki listelerini İK listesiyle karşılaştırır; ayrılan personel, sahipsiz hesap, aşırı yetki ve görev ayrılığı ihlallerini bulur.")
    p.add_argument("--personel", type=Path, default=ORNEK / "personel.csv", help="Sicil, Ad Soyad, E-posta, Departman, Pozisyon, Durum, Ayrılış Tarihi, Yönetici")
    p.add_argument("--yetkiler", type=Path, default=ORNEK / "yetkiler.csv",
                   help="Sistem, Kullanıcı Adı, Ad Soyad, E-posta, Sicil, Rol, Son Giriş, Hesap Durumu, Hesap Türü (her rol bir satır)")
    p.add_argument("--rol-matrisi", type=Path, help="Pozisyon, Sistem, İzinli Roller (| ile); Pozisyon '*' = herkes")
    p.add_argument("--gorev-ayriligi", type=Path, help="Kural, Rol A, Rol B, Risk (rol adlarında * joker)")
    p.add_argument("--kritik-roller", type=Path, help="Rol (her satır bir rol; * joker)")
    p.add_argument("--bugun", help="Gözden geçirme tarihi GG.AA.YYYY (örnek veride 09.10.2026)")
    p.add_argument("--pasif-gun", type=int, default=90, help="Bu kadar gün giriş yapılmayan aktif hesap işaretlenir (varsayılan 90)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "yetki_gozden_gecirme.xlsx")
    a = p.parse_args(argv)
    ornek = a.personel == ORNEK / "personel.csv"
    mat = a.rol_matrisi or (ORNEK / "rol_matrisi.csv" if ornek else None)
    sod = a.gorev_ayriligi or (ORNEK / "gorev_ayriligi.csv" if ornek else None)
    kri = a.kritik_roller or (ORNEK / "kritik_roller.csv" if ornek else None)
    for y in (a.personel, a.yetkiler, mat, sod, kri):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    bugun = tarih(a.bugun) if a.bugun else (date(2026, 10, 9) if ornek else date.today())
    if not bugun:
        print("[X] --bugun GG.AA.YYYY biçiminde olmalı")
        return 2
    try:
        s = calistir(a.personel, a.yetkiler, a.cikti, bugun, mat, sod, kri, a.pasif_gun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(s['personel'])} personel · {len(s['hesaplar'])} hesap · {len({h.sistem for h in s['hesaplar']})} sistem · "
          f"{len(s['ihlaller'])} görev ayrılığı ihlali")
    for b in s["bulgular"]:
        if b["onem"] == "Yüksek":
            print(f"[!] {b['sistem']} / {b['kullanici'] or b['kisi']} · {b['tur']}: {b['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
