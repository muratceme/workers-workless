import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 9)


def calistir(cikti, **kw):
    return main.calistir(ORNEK / "personel.csv", ORNEK / "yetkiler.csv", cikti, BUGUN, ORNEK / "rol_matrisi.csv", ORNEK / "gorev_ayriligi.csv",
                         ORNEK / "kritik_roller.csv", **kw)


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "y.xlsx"
        cls.s = calistir(cls.cikti)
        cls.b = {(b["sistem"], b["kullanici"], b["tur"]) for b in cls.s["bulgular"]}
        cls.h = {(h.sistem, h.kullanici): h for h in cls.s["hesaplar"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_eslestirme(self):
        self.assertEqual(self.h[("ERP", "mkaya")].eslesme, "Sicil")
        self.assertEqual(self.h[("ERP", "zdemir")].eslesme, "E-posta")
        self.assertEqual(self.h[("Banka Portalı", "AYILMAZ")].eslesme, "Ad soyad")
        self.assertEqual(self.h[("ERP", "acelik")].roller, ["Satın Alma Siparişi", "Tedarikçi Kartı Açma", "Mal Kabul"])
        self.assertEqual(len(self.s["hesaplar"]), 21)

    def test_ayrilan_ve_sahipsiz(self):
        self.assertTrue({("ERP", "zdemir", "Ayrılan personelin aktif hesabı"), ("Banka Portalı", "ZDEMIR", "Ayrılıştan sonra giriş"),
                         ("E-posta / AD", "burak.koc", "Ayrılıştan sonra giriş"), ("E-posta / AD", "murat.ekinci", "Sahipsiz hesap"),
                         ("ERP", "bkoc", "Ayrılan personelin pasif hesabı"), ("ERP", "stajyer01", "Ortak hesap"),
                         ("ERP", "entegrasyon", "Servis hesabı"), ("E-posta / AD", "elif.aksoy", "Kullanılmayan hesap")} <= self.b)
        self.assertNotIn(("Banka Portalı", "ZDEMIR", "Matris dışı yetki"), self.b)

    def test_matris_ve_sod(self):
        self.assertIn(("ERP", "mkaya", "Matris dışı yetki"), self.b)
        self.assertNotIn(("E-posta / AD", "deniz.aydin", "Matris dışı yetki"), self.b)     # '*' satırı
        kurallar = sorted((x["kisi"].ad, x["kural"]) for x in self.s["ihlaller"])
        self.assertEqual(kurallar, [("Ali Çelik", "Sipariş ve mal kabul"), ("Ali Çelik", "Tedarikçi açma ve sipariş"),
                                    ("Emre Öztürk", "Ödeme hazırlama ve onaylama"), ("Mehmet Kaya", "Muhasebe kaydı ve ödeme onayı")])
        self.assertEqual(self.h[("E-posta / AD", "cem.yildiz")].kritik, ["Domain Admin"])

    def test_excel(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Bulgular", "Yönetici Onay Listesi", "Görevler Ayrılığı", "Hesap Eşleşmesi", "Sistem Özeti"])
        yo = [r for r in wb["Yönetici Onay Listesi"].iter_rows(min_row=2, values_only=True)]
        self.assertTrue(any(r[0] == "Bilgi İşlem (sahipsiz)" and r[4] == "murat.ekinci" for r in yo))
        self.assertFalse(any(r[4] == "bkoc" for r in yo))                                  # kilitli hesap listede yok


class KuralTesti(unittest.TestCase):
    def test_pasif_gun_ve_joker(self):
        self.assertTrue(main.rol_uyar("Ödeme Onayı*", "Ödeme Onayı (2. imza)"))
        self.assertFalse(main.rol_uyar("Ödeme Onayı*", "Ödeme Hazırlama"))
        with tempfile.TemporaryDirectory() as tmp:
            s = calistir(Path(tmp) / "o.xlsx", pasif_gun=5)
            n = sum(1 for b in s["bulgular"] if b["tur"] == "Kullanılmayan hesap")
            self.assertGreater(n, 2)

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--bugun", "bugün"]), 2)
            self.assertEqual(main.main(["--yetkiler", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
