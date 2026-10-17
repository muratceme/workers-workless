"""
Vergi Takvimi Takibi — Gerisi Bende kod bloğu
Muhasebe › Muhasebe Müdürü
Şirketlerin beyan ve bildirim yükümlülüklerinden yaklaşan son tarihleri çıkarır, verilenleri işaretler, geciken ve yaklaşanları hatırlatır.
  KDV                         = dönemi izleyen ayın 28'i
  Muhtasar ve prim hizmet     = dönemi izleyen ayın 26'sı (aylık)
  Damga vergisi / konaklama   = dönemi izleyen ayın 26'sı
  Ba-Bs formları              = dönemi izleyen ayın son günü; Eylül 2024 döneminden itibaren kaldırıldı (565 Sıra No'lu VUK GT)
  Geçici vergi                = çeyreği izleyen 2. ayın 17'si; 4. dönem 2022-2024'te yoktu, 2025'ten itibaren 7566 sayılı Kanunla yeniden var
  Kurumlar vergisi            = hesap dönemini (takvim yılı) izleyen 30 Nisan
  Gelir vergisi               = izleyen yılın 31 Mart'ı (1. taksit); 2. taksit 31 Temmuz
  e-Defter berat (aylık)      = dönemi izleyen 4. ayın 10'u (gelir) / 14'ü (kurumlar); aralık: GV 10 Nisan, KV 14 Mayıs
  Son gün hafta sonu ya da resmî tatile rastlarsa izleyen ilk iş günü; GİB uzatmaları --uzatma ile girilir.
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
BABS_SON = (2024, 9)  # 565 Sıra No'lu VUK Genel Tebliği: Eylül 2024 döneminden itibaren Form Ba / Bs verilmez
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
YUK = {"sirket": ("sirket", "mukellef", "unvan", "firma"), "yuk": ("yukumluluk", "beyanname", "beyan", "bildirim"), "tur": ("vergi turu", "mukellefiyet"), "sorumlu": ("sorumlu", "personel")}
VERILEN = {"sirket": ("sirket", "mukellef", "unvan", "firma"), "yuk": ("yukumluluk", "beyanname", "beyan", "bildirim"), "donem": ("donem",), "tarih": ("verilis tarihi", "tarih", "onay tarihi")}
SABIT_TATIL = ((1, 1), (4, 23), (5, 1), (5, 19), (7, 15), (8, 30), (10, 29))
DINI_TATIL = {date(2026, 3, 20), date(2026, 3, 21), date(2026, 3, 22), date(2026, 5, 27), date(2026, 5, 28), date(2026, 5, 29), date(2026, 5, 30)}
TURLER = ("KDV", "Muhtasar ve prim hizmet", "Damga vergisi", "Konaklama vergisi", "Ba-Bs", "Geçici vergi", "Kurumlar vergisi", "Gelir vergisi", "e-Defter berat")


def tatil_mi(g: date, ek: set[date]) -> bool:
    return g.weekday() >= 5 or (g.month, g.day) in SABIT_TATIL or g in DINI_TATIL or g in ek


def is_gunu(g: date, ek: set[date]) -> date:
    while tatil_mi(g, ek):
        g += timedelta(days=1)
    return g


def _ay(y: int, a: int, ileri: int) -> tuple[int, int]:
    k, m = divmod(a - 1 + ileri, 12)
    return y + k, m + 1


def _ay_sonu(y: int, a: int) -> date:
    y2, a2 = _ay(y, a, 1)
    return date(y2, a2, 1) - timedelta(days=1)


def yukumluluk(x) -> str | None:
    k = katla(x)
    if k.startswith("kdv") or "katma deger" in k:
        return "KDV"
    if "muhtasar" in k or "prim hizmet" in k or k == "mphb":
        return "Muhtasar ve prim hizmet"
    if "damga" in k:
        return "Damga vergisi"
    if "konaklama" in k:
        return "Konaklama vergisi"
    if k.replace(" ", "") in ("babs", "formbabs", "baveyabs") or k.startswith("ba bs"):
        return "Ba-Bs"
    if "gecici" in k:
        return "Geçici vergi"
    if "kurumlar" in k:
        return "Kurumlar vergisi"
    if k.startswith("gelir vergisi") or k in ("yillik gelir", "gelir"):
        return "Gelir vergisi"
    if "defter" in k or "berat" in k:
        return "e-Defter berat"
    return None


def donemler(yuk: str, vt: str, yil: int) -> list[tuple[str, date, str]]:
    """Bir takvim yılına ait dönemler için (dönem etiketi, kaydırılmamış son gün, açıklama) listesi."""
    l = []
    if yuk in ("KDV", "Muhtasar ve prim hizmet", "Damga vergisi", "Konaklama vergisi", "Ba-Bs"):
        for a in range(1, 13):
            if yuk == "Ba-Bs" and (yil, a) >= BABS_SON:
                break  # Form Ba / Bs Eylül 2024 döneminden itibaren verilmez
            y2, a2 = _ay(yil, a, 1)
            g = _ay_sonu(y2, a2) if yuk == "Ba-Bs" else date(y2, a2, 28 if yuk == "KDV" else 26)
            l.append((f"{yil}/{a:02d}", g, "Beyan ve ödeme" if yuk != "Ba-Bs" else "Form Ba / Bs"))
    elif yuk == "Geçici vergi":
        # 4. dönem 7338 sayılı Kanunla 2022-2024 dönemlerinde kaldırılmıştı; 7566 sayılı Kanunla 2025'ten itibaren yeniden var
        for c in ((1, 2, 3, 4) if yil >= 2025 or yil < 2022 else (1, 2, 3)):
            y2, a2 = _ay(yil, c * 3, 2)
            l.append((f"{yil}/{c}. dönem", date(y2, a2, 17), "Beyan ve ödeme"))
    elif yuk == "Kurumlar vergisi":
        l.append((f"{yil}", date(yil + 1, 4, 30), "Yıllık beyan ve ödeme"))
    elif yuk == "Gelir vergisi":
        l += [(f"{yil}", date(yil + 1, 3, 31), "Yıllık beyan ve 1. taksit"), (f"{yil} · 2. taksit", date(yil + 1, 7, 31), "2. taksit ödemesi")]
    elif yuk == "e-Defter berat":
        for a in range(1, 13):
            gun = 10 if vt == "Gelir" else 14
            y2, a2 = _ay(yil, a, 5 if (a == 12 and vt != "Gelir") else 4)
            l.append((f"{yil}/{a:02d}", date(y2, a2, gun), "Oluşturma, imza ve berat yükleme"))
    return l


def _uzat(g: date, uzatma: dict) -> date:
    seen = set()
    while g in uzatma and g not in seen:
        seen.add(g)
        g = uzatma[g]
    return g


def yuk_oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, YUK, ("sirket", "yuk")):
        s = metin(r["sirket"])
        if not s:
            continue
        y = yukumluluk(r["yuk"])
        vt = katla(r["tur"])
        vt = "Gelir" if vt in ("gelir", "gelir vergisi", "gv") else "Kurumlar"
        if y is None:
            uy.append(uyari("Yüksek", "Yükümlülük tanınmadı", s, f"Satır {r['_satir']}: '{metin(r['yuk'])}' tanınmadı. Tanınanlar: {', '.join(TURLER)}."))
            continue
        if (katla(s), y) in gorulen:
            continue
        if y == "Ba-Bs" and not any(u["tur"] == "Ba-Bs kaldırıldı" for u in uy):
            uy.append(uyari("Bilgi", "Ba-Bs kaldırıldı", s, "Form Ba / Bs, 565 Sıra No'lu VUK Genel Tebliği ile Eylül 2024 döneminden itibaren kaldırıldı; yalnız Ağustos 2024 ve öncesi dönemler listelenir."))
        gorulen.add((katla(s), y))
        l.append({"sirket": s, "yuk": y, "vt": vt, "sorumlu": metin(r["sorumlu"])})
    return l, uy


def verilen_oku(yol: Path) -> tuple[dict, list[dict]]:
    uy, d = [], {}
    for r in kayitlar(yol, VERILEN, ("sirket", "yuk", "donem")):
        s = metin(r["sirket"])
        if not s:
            continue
        y = yukumluluk(r["yuk"])
        t = tarih(r["tarih"]) if metin(r["tarih"]) else None
        if y is None or (metin(r["tarih"]) and t is None) or not metin(r["donem"]):
            uy.append(uyari("Orta", "Verilen beyan satırı okunamadı", s, f"Satır {r['_satir']}: yükümlülük, dönem ya da tarih okunamadı."))
            continue
        d[(katla(s), y, _donem_anahtar(r["donem"]))] = t or date.min
    return d, uy


def _donem_anahtar(s) -> tuple:
    """'2026/09', '2026-9', '09.2026' → (2026, 9); '2026/2. dönem' → (2026, 2); '2025' → (2025,)."""
    n = [int(x) for x in re.findall(r"\d+", metin(s))]
    if len(n) == 2 and n[1] > 31 >= n[0]:
        n = [n[1], n[0]]
    return tuple(n)


def takvim(yukler: list[dict], verilen: dict, bugun: date, ileri: int = 45, geri: int = 60, ek: set[date] = frozenset(), uzatma: dict | None = None,
           uyari_gun: int = 5) -> dict:
    if not yukler:
        raise ValueError("Okunabilen yükümlülük yok.")
    uzatma = uzatma or {}
    bas, son = bugun - timedelta(days=geri), bugun + timedelta(days=ileri)
    uy, satir = [], []
    for y in yukler:
        for yil in range(bas.year - 2, son.year + 1):
            for dn, ham, ac in donemler(y["yuk"], y["vt"], yil):
                g = _uzat(ham if ham in uzatma else is_gunu(ham, ek), uzatma)
                if not bas <= g <= son:
                    continue
                v = verilen.get((katla(y["sirket"]), y["yuk"], _donem_anahtar(dn)))
                kalan = (g - bugun).days
                if v is not None:
                    durum = "Verildi" if v == date.min or v <= g else "Geç verildi"
                elif kalan < 0:
                    durum = "Gecikti"
                elif kalan == 0:
                    durum = "Bugün"
                elif kalan <= uyari_gun:
                    durum = "Yaklaşıyor"
                else:
                    durum = "Planlı"
                s = {"sirket": y["sirket"], "yuk": y["yuk"], "donem": dn, "aciklama": ac, "ham": ham, "son": g, "kaydi": g != ham, "kalan": kalan, "durum": durum,
                     "verilis": v if v and v != date.min else None, "sorumlu": y["sorumlu"]}
                satir.append(s)
                kim = f"{y['sirket']} · {y['yuk']} {dn}"
                if durum == "Gecikti":
                    uy.append(uyari("Yüksek", "Son tarih geçti", kim, f"Son gün {g:%d.%m.%Y} ({-kalan} gün önce); verildiğine dair kayıt yok."))
                elif durum in ("Bugün", "Yaklaşıyor"):
                    uy.append(uyari("Orta", "Son tarih yaklaşıyor", kim, f"Son gün {g:%d.%m.%Y}" + (" (bugün)" if kalan == 0 else f" ({kalan} gün)") + (f" · {y['sorumlu']}" if y["sorumlu"] else "") + "."))
                elif durum == "Geç verildi":
                    uy.append(uyari("Bilgi", "Geç verilmiş", kim, f"Verildi {v:%d.%m.%Y}, son gün {g:%d.%m.%Y}."))
    satir.sort(key=lambda s: (s["son"], TURLER.index(s["yuk"]), katla(s["sirket"])))
    ozet = {d: sum(1 for s in satir if s["durum"] == d) for d in ("Gecikti", "Bugün", "Yaklaşıyor", "Planlı", "Verildi", "Geç verildi")}
    return {"satirlar": satir, "ozet": ozet, "bas": bas, "son": son, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Vergi Takvimi"
    _baslik(ws, ["Son Gün", "Gün", "Şirket", "Yükümlülük", "Dönem", "İşlem", "Kalan Gün", "Durum", "Veriliş", "Sorumlu", "Not"], (11, 6, 26, 22, 14, 26, 8, 11, 11, 14, 30), "C2")
    gunad = ("Pzt", "Sal", "Çar", "Per", "Cum", "Cmt", "Paz")
    renk = {"Gecikti": KIRMIZI, "Bugün": SARI, "Yaklaşıyor": SARI, "Verildi": YESIL, "Geç verildi": MAVI}
    for s in r["satirlar"]:
        _satir(ws, [s["son"], gunad[s["son"].weekday()], s["sirket"], s["yuk"], s["donem"], s["aciklama"], s["kalan"], s["durum"], s["verilis"], s["sorumlu"],
                    f"Asıl gün {s['ham']:%d.%m.%Y}; tatil / uzatma ile kaydı" if s["kaydi"] else ""], tarih_s=(1, 9), dolgu={8: renk.get(s["durum"])})
    ws.append([])
    ws.append([f"Pencere {r['bas']:%d.%m.%Y} – {r['son']:%d.%m.%Y}. Son günler genel kurala göre hesaplanmıştır; GİB'in süre uzatma duyurularını --uzatma ile girin."])
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def tatil_oku(yol: Path) -> set[date]:
    return {g for r in tablo_oku(yol) for c in r if (g := tarih(c))}


def calistir(yuk_yol: Path, ver_yol: Path | None, cikti: Path, bugun: date, ileri: int = 45, geri: int = 60, ek: set[date] = frozenset(), uzatma: dict | None = None,
             uyari_gun: int = 5) -> dict:
    yk, uy1 = yuk_oku(yuk_yol)
    vr, uy2 = verilen_oku(ver_yol) if ver_yol else ({}, [])
    r = takvim(yk, vr, bugun, ileri, geri, ek, uzatma, uyari_gun)
    r["uyarilar"] = sirala(uy1 + uy2 + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Beyan ve bildirim son tarihlerini şirket bazında listeler, gecikenleri ve yaklaşanları hatırlatır.")
    p.add_argument("--yukumlulukler", type=Path, help="Şirket, Yükümlülük (KDV / Muhtasar ve prim hizmet / Damga / Konaklama / Ba-Bs / Geçici vergi / Kurumlar / Gelir / e-Defter berat), Vergi Türü, Sorumlu")
    p.add_argument("--verilenler", type=Path, help="Şirket, Yükümlülük, Dönem (2026/09 · 2026/2. dönem · 2025), Veriliş Tarihi")
    p.add_argument("--tarih", help="Rapor günü GG.AA.YYYY (varsayılan bugün; örnek veride 14.10.2026)")
    p.add_argument("--ileri", type=int, default=45, help="Kaç gün ileriye bakılır (varsayılan 45)")
    p.add_argument("--geri", type=int, default=60, help="Kaç gün geriye bakılır (varsayılan 60)")
    p.add_argument("--uyari", type=int, default=5, help="Son güne kaç gün kala uyarılır (varsayılan 5)")
    p.add_argument("--tatiller", type=Path, help="Ek tatil günleri listesi (2026 dışındaki dini bayramlar, idari izinler)")
    p.add_argument("--uzatma", action="append", default=[], help="Süre uzatımı ESKİ=YENİ (GG.AA.YYYY=GG.AA.YYYY)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "vergi_takvimi.xlsx")
    a = p.parse_args(argv)
    ornek = a.yukumlulukler is None
    yk = a.yukumlulukler or ORNEK / "yukumlulukler.csv"
    vr = a.verilenler or (ORNEK / "verilenler.csv" if ornek else None)
    bugun = tarih(a.tarih) if a.tarih else (date(2026, 10, 14) if ornek else date.today())
    uzatma = {}
    for u in a.uzatma:
        e, _, y = u.partition("=")
        ge, gy = tarih(e.strip()), tarih(y.strip())
        if ge is None or gy is None or gy <= ge:
            print(f"[X] --uzatma '{u}' ESKİ=YENİ biçiminde olmalı")
            return 2
        uzatma[ge] = gy
    if bugun is None or not 1 <= a.ileri <= 400 or not 0 <= a.geri <= 400 or a.uyari < 0:
        print("[X] --tarih GG.AA.YYYY, --ileri 1-400, --geri 0-400, --uyari 0 ya da üstü olmalı")
        return 2
    for y in (yk, vr, a.tatiller):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(yk, vr, a.cikti, bugun, a.ileri, a.geri, tatil_oku(a.tatiller) if a.tatiller else set(), uzatma, a.uyari)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {len(r['satirlar'])} son gün · geciken {o['Gecikti']} · bugün {o['Bugün']} · yaklaşan {o['Yaklaşıyor']} · planlı {o['Planlı']} · verilen {o['Verildi'] + o['Geç verildi']}")
    for s in [s for s in r["satirlar"] if s["durum"] in ("Bugün", "Yaklaşıyor", "Planlı")][:8]:
        print(f"     {s['son']:%d.%m.%Y} · {s['sirket']} · {s['yuk']} {s['donem']}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Takvim: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
