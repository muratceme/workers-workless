import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 14)


def son(yuk, vt, yil, donem, ek=frozenset()):
    return next(main.is_gunu(g, ek) for d, g, _ in main.donemler(yuk, vt, yil) if d == donem)


class KuralTesti(unittest.TestCase):
    def test_son_gunler(self):
        self.assertEqual([son("KDV", "Kurumlar", 2026, "2026/09"), son("Muhtasar ve prim hizmet", "Kurumlar", 2026, "2026/09"), son("Ba-Bs", "Kurumlar", 2026, "2026/09")],
                         [date(2026, 10, 28), date(2026, 10, 26), date(2026, 11, 2)])
        # KDV 2026/10: 28.11.2026 cumartesi → 30.11.2026; geçici vergi 1. dönem 17.05.2026 pazar → 18.05.2026
        self.assertEqual([son("KDV", "Kurumlar", 2026, "2026/10"), son("Geçici vergi", "Kurumlar", 2026, "2026/1. dönem"), son("Geçici vergi", "Kurumlar", 2026, "2026/3. dönem")],
                         [date(2026, 11, 30), date(2026, 5, 18), date(2026, 11, 17)])
        self.assertEqual([d for d, _, _ in main.donemler("Geçici vergi", "Gelir", 2026)], ["2026/1. dönem", "2026/2. dönem", "2026/3. dönem"])
        self.assertEqual([son("Kurumlar vergisi", "Kurumlar", 2025, "2025"), son("Gelir vergisi", "Gelir", 2025, "2025"), son("Gelir vergisi", "Gelir", 2025, "2025 · 2. taksit")],
                         [date(2026, 4, 30), date(2026, 3, 31), date(2026, 7, 31)])
        self.assertEqual([son("e-Defter berat", "Gelir", 2026, "2026/06"), son("e-Defter berat", "Kurumlar", 2026, "2026/06"), son("e-Defter berat", "Kurumlar", 2025, "2025/12")],
                         [date(2026, 10, 12), date(2026, 10, 14), date(2026, 5, 14)])
        self.assertEqual(son("Damga vergisi", "Kurumlar", 2026, "2026/04", {date(2026, 5, 26)}), date(2026, 6, 1))

    def test_donusum(self):
        self.assertEqual([main.yukumluluk(x) for x in ("KDV 1", "MPHB", "Form Ba-Bs", "Geçici Vergi (KV)", "e-defter beratı", "Yıllık gelir", "ÖTV")],
                         ["KDV", "Muhtasar ve prim hizmet", "Ba-Bs", "Geçici vergi", "e-Defter berat", "Gelir vergisi", None])
        self.assertEqual([main._donem_anahtar(x) for x in ("2026/09", "2026-9", "09.2026", "2026/2. dönem", "2025")], [(2026, 9), (2026, 9), (2026, 9), (2026, 2), (2025,)])


class TakvimTesti(unittest.TestCase):
    def test_takvim(self):
        y = [{"sirket": "A", "yuk": "KDV", "vt": "Kurumlar", "sorumlu": "S"}, {"sirket": "A", "yuk": "Muhtasar ve prim hizmet", "vt": "Kurumlar", "sorumlu": ""}]
        v = {("a", "KDV", (2026, 8)): date(2026, 9, 30), ("a", "Muhtasar ve prim hizmet", (2026, 8)): date(2026, 9, 20)}
        r = main.takvim(y, v, date(2026, 10, 24), 10, 40, uyari_gun=5)
        self.assertEqual([(s["yuk"], s["donem"], s["son"], s["durum"]) for s in r["satirlar"]],
                         [("KDV", "2026/08", date(2026, 9, 28), "Geç verildi"), ("Muhtasar ve prim hizmet", "2026/08", date(2026, 9, 28), "Verildi"),
                          ("Muhtasar ve prim hizmet", "2026/09", date(2026, 10, 26), "Yaklaşıyor"), ("KDV", "2026/09", date(2026, 10, 28), "Yaklaşıyor")])
        self.assertEqual([s["kaydi"] for s in r["satirlar"]][:2], [False, True])
        r = main.takvim(y, {}, date(2026, 10, 27), 3, 2, uzatma={date(2026, 10, 26): date(2026, 11, 3)})
        self.assertEqual([(s["yuk"], s["son"], s["durum"]) for s in r["satirlar"]], [("KDV", date(2026, 10, 28), "Yaklaşıyor")])
        r = main.takvim(y[:1], {}, date(2026, 6, 1), 10, 1, uzatma={date(2026, 5, 28): date(2026, 6, 5)})
        self.assertEqual([(s["donem"], s["son"], s["kaydi"]) for s in r["satirlar"]], [("2026/04", date(2026, 6, 5), True)])
        r = main.takvim(y, {}, date(2026, 10, 28), 1, 3)
        self.assertEqual([s["durum"] for s in r["satirlar"]], ["Gecikti", "Bugün"])
        with self.assertRaises(ValueError):
            main.takvim([], {}, BUGUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "yukumlulukler.csv", ORNEK / "verilenler.csv", cikti, BUGUN)
            self.assertEqual((len(r["satirlar"]), r["ozet"]["Gecikti"], r["ozet"]["Bugün"], r["ozet"]["Planlı"], r["ozet"]["Verildi"] + r["ozet"]["Geç verildi"]), (51, 1, 1, 22, 27))
            self.assertEqual([u["tur"] for u in r["uyarilar"] if u["onem"] == "Yüksek"], ["Son tarih geçti", "Yükümlülük tanınmadı"])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Vergi Takvimi", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            yy, vy, ty = Path(tmp) / "y.csv", Path(tmp) / "v.csv", Path(tmp) / "t.csv"
            yy.write_text("Firma;Beyanname;Mükellefiyet\nA;KDV;\nA;kdv;\nA;x;\nB;Geçici;GV\n", encoding="utf-8")
            vy.write_text("Firma;Beyanname;Dönem;Tarih\nA;KDV;2026/09;20.10.2026\nA;x;2026/09;\nA;KDV;;\n", encoding="utf-8")
            ty.write_text("28.10.2026\n", encoding="utf-8")
            l, uy = main.yuk_oku(yy)
            self.assertEqual(([(x["sirket"], x["yuk"], x["vt"]) for x in l], len(uy)), ([("A", "KDV", "Kurumlar"), ("B", "Geçici vergi", "Gelir")], 1))
            d, uy = main.verilen_oku(vy)
            self.assertEqual((d, len(uy)), ({("a", "KDV", (2026, 9)): date(2026, 10, 20)}, 2))
            self.assertEqual(main.tatil_oku(ty), {date(2026, 10, 28)})
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--yukumlulukler", str(yy), "--verilenler", str(vy), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "14.10.2026", "--ileri", "60", "--geri", "10", "--uyari", "3", "--tatiller", str(ty), "--uzatma", "26.10.2026=02.11.2026"]), 0)
            self.assertEqual(main.main(["--yukumlulukler", str(yy), "--cikti", c]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--ileri", "0"], ["--geri", "-1"], ["--uzatma", "x=y"])], [2] * 4)
            self.assertEqual(main.main(["--yukumlulukler", str(Path(tmp) / "yok.csv")]), 1)
            yy.write_text("Firma;Beyanname\nA;x\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
