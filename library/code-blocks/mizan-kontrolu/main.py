"""
Mizan Kontrolü — Workers / Workless kod bloğu
Muhasebe › Genel Muhasebe Uzmanı

Aylık mizanı Tekdüzen Hesap Planı kurallarına göre kontrol eder:
  - Biçim: borç = alacak denkliği, satır aritmetiği (borç − alacak = bakiye), aynı satırda hem borç hem alacak
    bakiyesi, ana hesap ile alt hesaplar toplamının tutması.
  - Ters bakiye: hesabın doğasına aykırı bakiye (düzenleyici "(-)" hesaplar dahil); ana hesap normal görünse de
    ters bakiyeli alt hesaplar (müşteriden alınan avans, satıcıya verilen avans, kredili mevduat) ve virman önerisi.
  - Mantık kontrolleri: kasa şişkinliği, ortaklardan alacaklar (adat), KDV hesaplarının mahsup edilmemesi,
    7/A yansıtma hesaplarının kapanmaması, birikmiş amortismanın varlık değerini aşması, karşılığın alacağı aşması.
  - Önceki aya göre: olağandışı değişim, yön değiştiren / yeni açılan / kapanan hesaplar, kümülatif gelir-gider
    hesabının azalması (iptal/ters kayıt belirtisi).
İnternete bağlanmaz.

Kullanım:
    python main.py                                         # örnek mizanlarla dener
    python main.py --mizan eylul.xlsx --onceki agustos.xlsx
    python main.py --mizan eylul.xlsx --esik 50000 --esik-oran 30 --kasa-orani 2
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
SIFIR = Decimal(0)
KURUS = Decimal("0.01")

# Tekdüzen Hesap Planı: adı "(-)" ile biten düzenleyici hesaplar
DUZENLEYICI_AKTIF = {"103", "119", "122", "124", "129", "137", "139", "158", "199", "222", "224", "229", "237", "239",
                     "241", "243", "244", "246", "247", "249", "257", "268", "278", "298", "299"}
DUZENLEYICI_PASIF = {"302", "308", "322", "337", "371", "402", "408", "422", "437", "501", "503", "580", "591"}

# Ters bakiyede sık görülen durumlar için açıklama ve öneri
TERS_ACIKLAMA = {
    "100": ("Kasa alacak bakiye veremez: kasada eksi para olamaz", "Belgesiz ödeme/kayıt hatası; vergi incelemesinde ciddi bulgu. Kayıtları hemen inceleyin."),
    "102": ("Banka hesabı alacak bakiyede: kredili mevduat (KMH) kullanımı", "Bilançoda 300 Banka Kredileri'ne virman yapın."),
    "12": ("Müşteri hesabı alacak bakiyede: müşteriden avans alınmış veya fazla tahsilat", "Dönem sonunda 340 Alınan Sipariş Avansları'na virman yapın veya iade edin."),
    "15": ("Stok hesabı alacak bakiyede: alış kaydı satıştan sonra işlenmiş veya maliyet hatalı", "Stok hareketlerini ve maliyet kayıtlarını kontrol edin."),
    "159": ("Verilen avans alacak bakiyede", "Avansın fatura ile kapatılması ters işlenmiş olabilir."),
    "191": ("İndirilecek KDV alacak bakiyede", "KDV mahsup veya iade kaydını kontrol edin."),
    "25": ("Maddi duran varlık alacak bakiyede: satış/çıkış kaydı maliyetten fazla", "Çıkış kaydını ve birikmiş amortisman kapatmasını kontrol edin."),
    "32": ("Satıcı hesabı borç bakiyede: satıcıya avans verilmiş veya fazla ödeme", "Dönem sonunda 159 Verilen Sipariş Avansları'na virman yapın veya iadesini isteyin."),
    "335": ("Personele borçlar borç bakiyede: personele avans verilmiş", "196 Personel Avansları'na virman yapın."),
    "340": ("Alınan avans borç bakiyede", "Avansın fatura ile kapatılması ters işlenmiş olabilir."),
    "360": ("Ödenecek vergi borç bakiyede: fazla veya mükerrer ödeme", "Tahakkuk ve ödeme kayıtlarını eşleştirin; fazla ödeme varsa mahsup/iade isteyin."),
    "361": ("Ödenecek SGK primi borç bakiyede: fazla veya mükerrer ödeme", "Tahakkuk ve ödeme kayıtlarını eşleştirin."),
    "391": ("Hesaplanan KDV borç bakiyede", "KDV mahsup kaydı fazla yapılmış olabilir."),
    "60": ("Satış hesabı borç bakiyede: iade/iskonto 61x yerine satış hesabına işlenmiş olabilir", "İade ve iskontoları 610/611 hesaplarında izleyin."),
}


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def katla(s) -> str:
    """Başlık karşılaştırması için Türkçe karakter ve ı/i farkını yok sayar."""
    s = kucuk(s)
    for a, b in zip("ıöüşğç", "iousgc"):
        s = s.replace(a, b)
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def para(x) -> Decimal:
    if x in (None, ""):
        return SIFIR
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("₺", "").replace(" ", "")
    eksi = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1:
        s = s.replace(".", "")
    try:
        d = Decimal(s)
    except InvalidOperation:
        return SIFIR
    return -d if eksi else d


def tl(x) -> str:
    return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


# ----------------------------------------------------------------------------
# Mizan okuma
# ----------------------------------------------------------------------------

def mizan_oku(yol: Path) -> dict[str, dict]:
    """Dönüş: kod → {ad, borc, alacak, bb, ab, net, satir}. net = borç bakiye − alacak bakiye."""
    s = tablo_oku(yol)
    bi = next((i for i, r in enumerate(s[:15]) if any(katla(c) in ("hesap kodu", "hesap no", "kod") for c in r)), None)
    if bi is None:
        raise SystemExit(f"{yol.name}: 'Hesap Kodu' başlıklı satır bulunamadı.")
    b = [katla(x) for x in s[bi]]

    def bul(*adlar):
        return next((b.index(katla(a)) for a in adlar if katla(a) in b), None)

    i_k, i_ad = bul("hesap kodu", "hesap no", "kod"), bul("hesap adı", "hesap ismi", "açıklama", "ad")
    i_b, i_a = bul("borç", "borç toplamı", "toplam borç"), bul("alacak", "alacak toplamı", "toplam alacak")
    i_bb = bul("borç bakiye", "borç bakiyesi", "bakiye borç")
    i_ab = bul("alacak bakiye", "alacak bakiyesi", "bakiye alacak")
    i_bak = bul("bakiye", "net bakiye")
    if not ((i_bb is not None and i_ab is not None) or (i_b is not None and i_a is not None) or i_bak is not None):
        raise SystemExit(f"{yol.name}: Borç/Alacak Bakiye (veya Borç/Alacak ya da Bakiye) sütunları gerekli. Başlıklar: {s[bi]}")
    hesaplar: dict[str, dict] = {}
    for n, r in enumerate(s[bi + 1:], bi + 2):
        kod = re.sub(r"\s", "", str(r[i_k] if i_k < len(r) else "") or "")
        if not re.match(r"^\d{3}", kod):
            continue                                      # toplam/ara başlık satırları
        al = lambda i: r[i] if i is not None and i < len(r) else None  # noqa: E731
        h = {"kod": kod, "ad": str(al(i_ad) or "").strip(), "satir": n, "borc": None, "alacak": None, "bb": None, "ab": None}
        if i_b is not None and i_a is not None:
            h["borc"], h["alacak"] = para(al(i_b)), para(al(i_a))
        if i_bb is not None and i_ab is not None:
            h["bb"], h["ab"] = para(al(i_bb)), para(al(i_ab))
            h["net"] = h["bb"] - h["ab"]
        elif h["borc"] is not None:
            h["net"] = h["borc"] - h["alacak"]
        else:
            h["net"] = para(al(i_bak))
        if kod in hesaplar:                               # aynı kod iki kez: topla
            for k in ("borc", "alacak", "bb", "ab", "net"):
                if h[k] is not None:
                    hesaplar[kod][k] = (hesaplar[kod][k] or SIFIR) + h[k]
        else:
            hesaplar[kod] = h
    if not hesaplar:
        raise SystemExit(f"{yol.name}: hesap satırı bulunamadı.")
    return hesaplar


def ust_kod(kod: str, kodlar: set[str]) -> str | None:
    """Mizanda bulunan en yakın üst hesap (120.01.001 → 120.01 → 120)."""
    parcalar = re.split(r"[.\-_/ ]", kod)
    for i in range(len(parcalar) - 1, 0, -1):
        aday = kod[:len(".".join(parcalar[:i]))]
        if aday in kodlar:
            return aday
    if len(kod) > 3 and kod[:3] in kodlar and kod != kod[:3]:
        return kod[:3]
    return None


def ana_hesaplar(h: dict[str, dict]) -> dict[str, dict]:
    """3 haneli ana hesap bakiyeleri: ana hesap satırı varsa o, yoksa en alt kırılımların toplamı."""
    kodlar = set(h)
    cocuklu = {ust_kod(k, kodlar) for k in kodlar} - {None}
    ana: dict[str, dict] = {}
    for kod in sorted(h):
        a = kod[:3]
        if kod == a:
            ana[a] = {"kod": a, "ad": h[kod]["ad"], "net": h[kod]["net"], "borc": h[kod]["borc"], "alacak": h[kod]["alacak"]}
        elif a not in h and kod not in cocuklu:
            x = ana.setdefault(a, {"kod": a, "ad": "", "net": SIFIR, "borc": None, "alacak": None})
            x["net"] += h[kod]["net"]
            for k in ("borc", "alacak"):
                if h[kod][k] is not None:
                    x[k] = (x[k] or SIFIR) + h[kod][k]
    return ana


# ----------------------------------------------------------------------------
# Hesap doğası
# ----------------------------------------------------------------------------

def doga(kod: str) -> str | None:
    """'B' borç bakiye, 'A' alacak bakiye, None: iki yönlü olabilir / kontrol edilmez."""
    k = kod[:3]
    if not k.isdigit():
        return None
    if k in DUZENLEYICI_AKTIF:
        return "A"
    if k in DUZENLEYICI_PASIF:
        return "B"
    s = k[0]
    if s in "12":
        return "B"
    if s in "345":
        return "A"
    if s == "6":
        if k[:2] in ("60", "64", "67"):
            return "A"
        if k[:2] in ("61", "62", "63", "65", "66", "68"):
            return "B"
        return "B" if k == "691" else None              # 690/692/697/698 kâr veya zarar
    if s == "7":
        if k[:2] in ("71", "72", "73", "74", "75", "76", "77", "78"):
            return {"0": "B", "1": "A"}.get(k[2])          # x0 gider, x1 yansıtma, diğerleri fark hesabı
        if k[:2] == "79":
            return "A" if k == "798" else ("B" if k <= "797" else None)
    return None


def ters_aciklama(kod: str) -> tuple[str, str]:
    for anahtar in (kod[:3], kod[:2]):
        if anahtar in TERS_ACIKLAMA:
            return TERS_ACIKLAMA[anahtar]
    return ("Hesabın doğasına aykırı bakiye", "Son kayıtları ve hesap kodu seçimini kontrol edin.")


# ----------------------------------------------------------------------------
# Kontroller
# ----------------------------------------------------------------------------

def bulgu(liste, seviye, kod, ad, tutar, kontrol, aciklama, oneri="", onceki=None):
    liste.append({"seviye": seviye, "kod": kod, "ad": ad, "tutar": tutar, "onceki": onceki, "kontrol": kontrol,
                  "aciklama": aciklama, "oneri": oneri})


def bicim_kontrolleri(h: dict, ana: dict, tolerans: Decimal) -> tuple[list, dict]:
    b = []
    # Satır aritmetiği
    for x in h.values():
        if x["bb"] is not None and x["bb"] > 0 and x["ab"] > 0:
            bulgu(b, "Hata", x["kod"], x["ad"], x["net"], "Satır", "Aynı satırda hem borç hem alacak bakiyesi var",
                  "Mizan raporunun bakiye sütunlarını kontrol edin.")
        if x["borc"] is not None and x["bb"] is not None and abs((x["borc"] - x["alacak"]) - (x["bb"] - x["ab"])) > tolerans:
            bulgu(b, "Hata", x["kod"], x["ad"], x["net"], "Satır aritmetiği",
                  f"Borç − alacak ({tl(x['borc'] - x['alacak'])}) bakiyeye ({tl(x['net'])}) eşit değil (satır {x['satir']})",
                  "Mizan elle düzenlenmiş veya hatalı aktarılmış olabilir.")
    # Ana hesap = alt hesaplar toplamı (yalnız doğrudan alt hesaplar)
    kodlar = set(h)
    cocuk = defaultdict(list)
    for k in kodlar:
        u = ust_kod(k, kodlar)
        if u:
            cocuk[u].append(k)
    for u, cs in sorted(cocuk.items()):
        toplam = sum((h[c]["net"] for c in cs), SIFIR)
        if abs(toplam - h[u]["net"]) > tolerans:
            bulgu(b, "Hata", u, h[u]["ad"], h[u]["net"], "Alt hesap toplamı",
                  f"Alt hesaplar toplamı {tl(toplam)}, hesap bakiyesi {tl(h[u]['net'])} (fark {tl(h[u]['net'] - toplam)})",
                  "Eksik aktarılmış alt hesap veya doğrudan ana hesaba yapılmış kayıt olabilir.")
    # Denklik (ana hesaplar üzerinden; 0-7 grupları; nazım 9 ayrı denkleşir)
    net = sum((a["net"] for k, a in ana.items() if k[0] in "01234567"), SIFIR)
    hareket = None
    if all(a["borc"] is not None for a in ana.values()):
        hb = sum((a["borc"] for a in ana.values()), SIFIR)
        ha = sum((a["alacak"] for a in ana.values()), SIFIR)
        hareket = (hb, ha)
        if abs(hb - ha) > tolerans:
            bulgu(b, "Hata", "-", "Mizan geneli", hb - ha, "Denklik", f"Borç toplamı {tl(hb)} ≠ alacak toplamı {tl(ha)}",
                  "Eksik hesap veya filtrelenmiş mizan olabilir.")
    if abs(net) > tolerans:
        bulgu(b, "Hata", "-", "Mizan geneli", net, "Denklik",
              f"Borç bakiyeler − alacak bakiyeler = {tl(net)}; mizan denk değil", "Eksik hesap veya filtrelenmiş mizan olabilir.")
    return b, {"net": net, "hareket": hareket}


def ters_bakiye_kontrolleri(h: dict, ana: dict, tolerans: Decimal) -> list:
    b = []
    ters_ana = set()
    for k, a in sorted(ana.items()):
        d = doga(k)
        if d and ((d == "B" and a["net"] < -tolerans) or (d == "A" and a["net"] > tolerans)):
            ters_ana.add(k)
            ac, on = ters_aciklama(k)
            bulgu(b, "Yüksek" if k in ("100", "191", "391") or k[0] in "2" else "Orta", k, a["ad"] or h.get(k, {}).get("ad", ""),
                  a["net"], "Ters bakiye", f"{ac}. Beklenen: {'borç' if d == 'B' else 'alacak'} bakiye", on)
    kodlar = set(h)
    cocuklu = {ust_kod(k, kodlar) for k in kodlar} - {None}
    for k, x in sorted(h.items()):
        if k == k[:3] or k in cocuklu or k[:3] in ters_ana:
            continue                                      # yalnız en alt kırılım; ana hesabı zaten terse düşenler hariç
        d = doga(k)
        if d and ((d == "B" and x["net"] < -tolerans) or (d == "A" and x["net"] > tolerans)):
            ac, on = ters_aciklama(k)
            bulgu(b, "Orta" if k[:3] in ("100", "102") else "Bilgi", k, x["ad"], x["net"], "Alt hesap ters bakiye", ac, on)
    return b


def mantik_kontrolleri(ana: dict, h: dict, kasa_orani: float, tolerans: Decimal) -> list:
    b = []
    n = lambda k: ana[k]["net"] if k in ana else SIFIR  # noqa: E731
    ad = lambda k: (ana[k]["ad"] if k in ana else "") or h.get(k, {}).get("ad", "")  # noqa: E731
    aktif = sum((a["net"] for k, a in ana.items() if k[0] in "12"), SIFIR)
    if n("100") > 0 and aktif > 0 and n("100") / aktif * 100 > Decimal(str(kasa_orani)):
        bulgu(b, "Orta", "100", ad("100"), n("100"), "Kasa şişkinliği",
              f"Kasa / aktif toplamı: %{float(n('100') / aktif * 100):.1f} (eşik %{kasa_orani:g})".replace(".", ","),
              "Fiili kasa sayımıyla karşılaştırın. Gerçekte olmayan kasa bakiyesi incelemede örtülü kazanç dağıtımı veya "
              "örtülü sermaye olarak değerlendirilebilir; adat hesaplanması istenebilir.")
    for k in ("131", "231"):
        if n(k) > tolerans:
            bulgu(b, "Orta", k, ad(k), n(k), "Ortaklardan alacaklar",
                  "Ortaklara kullandırılan tutar var", "Emsal faiz üzerinden adat hesaplayıp faiz faturası düzenleyin "
                  "(transfer fiyatlandırması yoluyla örtülü kazanç, KVK md. 13). Mali müşavirinize danışın.")
    if n("191") > tolerans and n("391") < -tolerans:
        bulgu(b, "Dikkat", "191/391", "İndirilecek / Hesaplanan KDV", n("191") + n("391"), "KDV mahsubu",
              f"191 ({tl(n('191'))}) ve 391 ({tl(-n('391'))}) birlikte bakiye veriyor",
              "Ay sonu KDV mahsup kaydı (191/391 → 190 veya 360) yapılmamış olabilir. Beyanname öncesi mizansa normaldir.")
    for g in ("71", "72", "73", "74", "75", "76", "77", "78"):
        gider, yans = n(g + "0"), n(g + "1")
        if (gider or yans) and abs(gider + yans) > tolerans:
            bulgu(b, "Dikkat", f"{g}0/{g}1", ad(g + "0") or "Maliyet hesabı", gider + yans, "Yansıtma",
                  f"{g}0 gider {tl(gider)}, {g}1 yansıtma {tl(-yans)}: fark {tl(gider + yans)}",
                  "7/A seçeneğinde ay sonunda gider hesabı yansıtma hesabıyla kapanmalı; yansıtma kaydını kontrol edin.")
    mdv = sum((n(k) for k in ("250", "251", "252", "253", "254", "255", "256")), SIFIR)
    if -n("257") > mdv + tolerans:
        bulgu(b, "Yüksek", "257", ad("257"), n("257"), "Amortisman",
              f"Birikmiş amortisman ({tl(-n('257'))}) maddi duran varlıklar toplamını ({tl(mdv)}) aşıyor",
              "Satılan/hurdaya çıkan varlıkların birikmiş amortismanı kapatılmamış olabilir.")
    for alacak, karsilik in (("128", "129"), ("228", "229"), ("138", "139")):
        if -n(karsilik) > n(alacak) + tolerans:
            bulgu(b, "Orta", karsilik, ad(karsilik), n(karsilik), "Karşılık",
                  f"Karşılık ({tl(-n(karsilik))}) şüpheli alacaktan ({tl(n(alacak))}) fazla",
                  "Tahsil edilen veya silinen şüpheli alacakların karşılığı iptal edilmemiş olabilir.")
    if n("590") and n("591"):
        bulgu(b, "Dikkat", "590/591", "Dönem net kârı / zararı", n("590") + n("591"), "Dönem sonucu",
              "Dönem net kârı ve dönem net zararı birlikte bakiye veriyor", "Kapanış kayıtlarını kontrol edin.")
    return b


def degisim_kontrolleri(ana: dict, onceki: dict, esik: Decimal, esik_oran: float, tolerans: Decimal) -> tuple[list, bool]:
    b = []
    satis = lambda m: sum((abs(a["net"]) for k, a in m.items() if k[:2] == "60"), SIFIR)  # noqa: E731
    yil_basi = satis(onceki) > 0 and satis(ana) < satis(onceki) / 2
    for k in sorted(set(ana) | set(onceki)):
        cur = ana[k]["net"] if k in ana else SIFIR
        pre = onceki[k]["net"] if k in onceki else SIFIR
        ad = (ana.get(k) or onceki.get(k))["ad"]
        fark = cur - pre
        if abs(cur) <= tolerans and abs(pre) <= tolerans:
            continue
        sonuc = k[0] in "67"
        if sonuc:
            if yil_basi:
                continue
            if abs(cur) + tolerans < abs(pre):
                bulgu(b, "Orta", k, ad, cur, "Kümülatif azalma",
                      f"Kümülatif gelir/gider hesabı önceki aya göre azalmış ({tl(pre)} → {tl(cur)})",
                      "Önceki ay kaydının iptali veya ters kayıt olabilir; nedenini belgeleyin.", pre)
            continue
        if k in ("190", "191", "391"):
            continue                                      # KDV hesapları her ay kapanıp yeniden açılır
        if abs(pre) <= tolerans:
            if abs(cur) >= esik:
                bulgu(b, "Bilgi", k, ad, cur, "Yeni bakiye", "Önceki ay bakiyesi olmayan hesapta önemli bakiye", "Kaydın doğruluğunu kontrol edin.", pre)
            continue
        if abs(cur) <= tolerans:
            if abs(pre) >= esik:
                bulgu(b, "Bilgi", k, ad, cur, "Kapanan hesap", "Önceki ay önemli bakiyesi olan hesap sıfırlanmış", "Kapanışın beklenen bir işlem olduğunu doğrulayın.", pre)
            continue
        d = doga(k)
        if d and ((d == "B" and cur < 0) or (d == "A" and cur > 0)):
            continue                                      # ters bakiye olarak zaten raporlandı
        if (cur > 0) != (pre > 0) and d is not None:
            bulgu(b, "Dikkat", k, ad, cur, "Yön değişimi", f"Bakiye yönü değişti ({tl(pre)} → {tl(cur)})", "Ters bakiyeye geçişin nedenini inceleyin.", pre)
            continue
        oran = abs(fark) / abs(pre) * 100
        if abs(fark) >= esik and oran >= Decimal(str(esik_oran)):
            bulgu(b, "Bilgi", k, ad, cur, "Olağandışı değişim",
                  f"Bakiye {tl(abs(pre))} → {tl(abs(cur))} ({'+' if abs(cur) > abs(pre) else '−'}%{float(oran):.0f})", "Değişimin nedenini açıklayın.", pre)
    return b, yil_basi


SEVIYE_SIRA = {"Hata": 0, "Yüksek": 1, "Orta": 2, "Dikkat": 3, "Bilgi": 4}


def calistir(mizan: Path, cikti: Path, onceki: Path | None = None, esik: float = 100000.0, esik_oran: float = 50.0,
             kasa_orani: float = 3.0, tolerans: float = 1.0) -> dict:
    tol = Decimal(str(tolerans))
    h = mizan_oku(mizan)
    ana = ana_hesaplar(h)
    bulgular, denklik = bicim_kontrolleri(h, ana, tol)
    bulgular += ters_bakiye_kontrolleri(h, ana, tol)
    bulgular += mantik_kontrolleri(ana, h, kasa_orani, tol)
    onceki_ana, yil_basi = None, False
    if onceki:
        onceki_ana = ana_hesaplar(mizan_oku(onceki))
        d, yil_basi = degisim_kontrolleri(ana, onceki_ana, Decimal(str(esik)), esik_oran, tol)
        bulgular += d
    bulgular.sort(key=lambda x: (SEVIYE_SIRA[x["seviye"]], x["kontrol"], x["kod"]))
    notlar = []
    if yil_basi:
        notlar.append("Satış hesapları önceki mizana göre yarıdan fazla azalmış: yeni yıl başı varsayıldı, "
                      "gelir-gider hesaplarında kümülatif karşılaştırma yapılmadı.")
    _rapor(h, ana, onceki_ana, bulgular, denklik, notlar, cikti, mizan.name, onceki.name if onceki else None)
    return {"hesaplar": h, "ana": ana, "bulgular": bulgular, "denklik": denklik, "notlar": notlar}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Hata": "F8C9C6", "Yüksek": "FDE2E1", "Orta": "FFE8CC", "Dikkat": "FFF4CE", "Bilgi": "E8F0FE"}
PARA = "#,##0.00;[Red]-#,##0.00"


def _baslik(ws, satir=1):
    for c in ws[satir]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI


def _rapor(h, ana, onceki_ana, bulgular, denklik, notlar, cikti, ad, onceki_ad):
    wb = Workbook()
    o = wb.active
    o.title = "Özet"
    o.append([f"Mizan kontrolü · {ad}" + (f" (önceki: {onceki_ad})" if onceki_ad else "")])
    o["A1"].font = Font(bold=True, size=12)
    o.append(["Hesap satırı", len(h)])
    o.append(["Ana hesap", len(ana)])
    if denklik["hareket"]:
        o.append(["Borç toplamı", float(denklik["hareket"][0])])
        o.append(["Alacak toplamı", float(denklik["hareket"][1])])
    o.append(["Bakiye farkı (borç − alacak)", float(denklik["net"])])
    for r in range(4, o.max_row + 1):
        o.cell(r, 2).number_format = PARA
    o.append([])
    o.append(["Seviye", "Bulgu sayısı"])
    _baslik(o, o.max_row)
    for s in SEVIYE_SIRA:
        o.append([s, sum(1 for x in bulgular if x["seviye"] == s)])
        o.cell(o.max_row, 1).fill = PatternFill("solid", fgColor=RENK[s])
    for nt in notlar:
        o.append(["Not", nt])
    o.append([])
    o.append(["Hata: mizan biçimi/aritmetiği bozuk · Yüksek: olmaması gereken bakiye · Orta: düzeltme/virman gerektirebilir · "
              "Dikkat: ay sonu işlemi eksik olabilir · Bilgi: açıklanması gereken değişim"])
    o.column_dimensions["A"].width = 34
    o.column_dimensions["B"].width = 20

    b = wb.create_sheet("Bulgular")
    b.append(["Seviye", "Kontrol", "Hesap Kodu", "Hesap Adı", "Bakiye (B+ / A−)", "Önceki Bakiye", "Açıklama", "Öneri", "Açıklama / Sorumlu"])
    _baslik(b)
    for x in bulgular:
        b.append([x["seviye"], x["kontrol"], x["kod"], x["ad"], float(x["tutar"]),
                  None if x["onceki"] is None else float(x["onceki"]), x["aciklama"], x["oneri"], None])
        b.cell(b.max_row, 1).fill = PatternFill("solid", fgColor=RENK[x["seviye"]])
        for c in (5, 6):
            b.cell(b.max_row, c).number_format = PARA
    for j, w in enumerate((9, 22, 13, 30, 16, 16, 70, 70, 24), 1):
        b.column_dimensions[get_column_letter(j)].width = w
    b.freeze_panes = "C2"
    b.auto_filter.ref = b.dimensions

    a = wb.create_sheet("Ana Hesaplar")
    a.append(["Hesap Kodu", "Hesap Adı", "Beklenen Bakiye", "Bakiye (B+ / A−)"] +
             (["Önceki", "Değişim", "Değişim %"] if onceki_ana is not None else []))
    _baslik(a)
    for k in sorted(set(ana) | set(onceki_ana or {})):
        cur = ana[k]["net"] if k in ana else SIFIR
        d = doga(k)
        satir = [k, (ana.get(k) or (onceki_ana or {}).get(k))["ad"], {"B": "Borç", "A": "Alacak"}.get(d, "İki yönlü"), float(cur)]
        if onceki_ana is not None:
            pre = onceki_ana[k]["net"] if k in onceki_ana else SIFIR
            satir += [float(pre), float(cur - pre), float((cur - pre) / abs(pre)) if pre else None]
        a.append(satir)
        for c in (4, 5, 6):
            a.cell(a.max_row, c).number_format = PARA
        a.cell(a.max_row, 7).number_format = "0.0%"
        if d and ((d == "B" and cur < 0) or (d == "A" and cur > 0)):
            a.cell(a.max_row, 4).fill = PatternFill("solid", fgColor=RENK["Yüksek"])
    for j, w in enumerate((11, 40, 14, 18, 18, 18, 11), 1):
        a.column_dimensions[get_column_letter(j)].width = w
    a.freeze_panes = "C2"

    bi = wb.create_sheet("Bilgi")
    for s in [["Bakiye işareti", "Borç bakiye pozitif (+), alacak bakiye negatif (−) gösterilir"],
              ["Ana hesap", "3 haneli hesap satırı varsa o kullanılır; yoksa en alt kırılımların toplamı"],
              ["Ters bakiye", "Tekdüzen Hesap Planı'na göre: 1-2 borç, 3-4-5 alacak; adı (-) ile biten düzenleyici hesaplar ters; "
                              "60/64/67 alacak, 61-63/65/66/68 borç; 7/A'da x0 borç, x1 alacak; 690/692/697/698 ve 8-9 grupları kontrol dışı"],
              ["Alt hesap", "Ana hesabı normal görünen ama en alt kırılımı ters bakiyeli hesaplar (müşteri/satıcı avansı, KMH) ayrıca listelenir"],
              ["Kümülatif", "6 ve 7 grubu hesaplar yıl başından itibaren birikir; önceki aya göre azalma iptal/ters kayıt belirtisidir"],
              ["Uyarı", "Bulgular kontrol önerisidir; kesin değerlendirme ve düzeltme kayıtları için mali müşavirinize danışın"]]:
        bi.append(s)
    bi.column_dimensions["A"].width = 16
    bi.column_dimensions["B"].width = 140
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ornek = BURASI / "ornek_veri"
    ap = argparse.ArgumentParser(description="Aylık mizanı ters bakiye, mantık ve önceki aya göre değişim açısından kontrol eder.")
    ap.add_argument("--mizan", type=Path, default=ornek / "mizan_eylul.csv", help="Kontrol edilecek mizan (.xlsx/.csv)")
    ap.add_argument("--onceki", type=Path, help="Önceki ay mizanı (isteğe bağlı)")
    ap.add_argument("--esik", type=float, default=100000.0, help="Değişim kontrolü tutar eşiği, TL (varsayılan 100000)")
    ap.add_argument("--esik-oran", type=float, default=50.0, help="Değişim kontrolü oran eşiği, %% (varsayılan 50)")
    ap.add_argument("--kasa-orani", type=float, default=3.0, help="Kasa / aktif toplamı uyarı eşiği, %% (varsayılan 3)")
    ap.add_argument("--tolerans", type=float, default=1.0, help="Yuvarlama toleransı, TL (varsayılan 1)")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "mizan_kontrolu.xlsx")
    a = ap.parse_args(argv)
    if a.mizan == ornek / "mizan_eylul.csv" and not a.onceki:
        a.onceki = ornek / "mizan_agustos.csv"
    s = calistir(a.mizan, a.cikti, a.onceki, a.esik, a.esik_oran, a.kasa_orani, a.tolerans)
    from collections import Counter
    sayac = Counter(x["seviye"] for x in s["bulgular"])
    print(f"[OK] {len(s['hesaplar'])} hesap satırı · bakiye farkı {tl(s['denklik']['net'])} · "
          + " · ".join(f"{k}: {sayac.get(k, 0)}" for k in SEVIYE_SIRA))
    for x in s["bulgular"]:
        if x["seviye"] in ("Hata", "Yüksek"):
            print(f"[!] {x['seviye']} · {x['kod']} {x['ad']}: {x['aciklama']}")
    for nt in s["notlar"]:
        print(f"[i] {nt}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
