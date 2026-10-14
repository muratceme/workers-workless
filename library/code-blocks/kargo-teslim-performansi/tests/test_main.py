import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
GUN = date(2026, 9, 30)


def g(no, firma, v, t, durum="Teslim edildi", il=""):
    return {"no": no, "firma": firma, "verilis": date(2026, 9, v), "teslim": date(2026, 9, t) if t else None, "durum": durum, "il": il}


def veri():
    return [g("1", "F1", 1, 3, il="A"), g("2", "f1", 1, 5, il="A"), g("3", "F1", 2, 4, "Hasarlı", "B"), g("4", "F1", 20, None, "Yolda", "B"), g("5", "F1", 1, None, "Kayıp"),
            g("6", "F2", 1, 2), g("7", "F2", 1, None, "İade"), g("8", "F2", 1, 3)]


def k(f):
    return (f["gonderi"], f["teslim"], f["ort"], f["medyan"], f["en_uzun"], f["zamaninda"], f["hasar_oran"], f["kayip_oran"], f["iade_oran"], f["acik"])


class AnalizTesti(unittest.TestCase):
    def test_gostergeler(self):
        r = main.analiz(veri(), GUN)
        self.assertEqual([(f["firma"],) + k(f) for f in r["firmalar"]],
                         [("F1", 5, 3, D("2.67"), D(2), 4, D("66.67"), D("20.00"), D("20.00"), D("0.00"), 1), ("F2", 3, 2, D("1.50"), D("1.5"), 2, D("100.00"), D("0.00"), D("0.00"), D("33.33"), 0)])
        self.assertEqual(k(r["toplam"]), (8, 5, D("2.20"), D(2), 4, D("80.00"), D("12.50"), D("12.50"), D("12.50"), 1))
        self.assertEqual([(f["firma"], f["il"], f["gonderi"], f["ort"]) for f in r["iller"]], [("F1", "A", 2, D("3.00")), ("F1", "B", 2, D("2.00"))])
        self.assertEqual([(x["no"], x["yas"]) for x in r["acik"]], [("4", 10)])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Orta", "Gecikmiş gönderi", "F1 · 4"), ("Orta", "Hasar oranı yüksek", "F1"), ("Orta", "Zamanında teslim düşük", "F1"), ("Yüksek", "Kayıp oranı yüksek", "F1")])

    def test_esikler(self):
        r = main.analiz(veri(), GUN, hedef=4, gecikme=10, zamaninda=D(100), hasar=D(20), kayip=D(20))
        self.assertEqual(([f["zamaninda"] for f in r["firmalar"]], r["acik"], r["uyarilar"]), ([D("100.00"), D("100.00")], [], []))
        self.assertIsNone(main.analiz([g("9", "F", 1, None, "Yolda")], GUN)["toplam"]["ort"])
        self.assertEqual([main._medyan(x) for x in ([], [5], [1, 9, 3], [4, 1, 2, 8])], [None, D(5), D(3), D(3)])
        with self.assertRaises(ValueError):
            main.analiz([], GUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "kargo_gonderileri.csv", cikti, GUN)
            self.assertEqual([(f["firma"],) + k(f) for f in r["firmalar"]],
                             [("Kargo A", 69, 67, D("2.45"), D(2), 5, D("80.60"), D("1.45"), D("0.00"), D("0.00"), 2), ("Kargo B", 56, 49, D("3.10"), D(3), 5, D("63.27"), D("1.79"), D("1.79"), D("5.36"), 3),
                              ("Kargo C", 55, 51, D("2.37"), D(2), 5, D("88.24"), D("0.00"), D("0.00"), D("3.64"), 2)])
            self.assertEqual([(x["no"], x["yas"]) for x in r["acik"]], [("KG70178", 12), ("KG70179", 9)])
            self.assertEqual([u["tur"] for u in r["uyarilar"]][:3], ["Kayıp oranı yüksek", "Gecikmiş gönderi", "Gecikmiş gönderi"])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Firmalar", "Firma ve İl", "Gecikmiş Gönderiler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "g.csv"
            y.write_text("Takip No;Taşıyıcı;Gönderim Tarihi;Teslimat Tarihi;Son Durum\n1;K;01.09.2026;03.09.2026;\n1;K;01.09.2026;;\n2;K;01.09.2026;;\n3;K;x;;\n4;K;05.09.2026;01.09.2026;\n"
                         "5;K;01.09.2026;;uçuyor\n6;K;01.09.2026;;Teslim edildi\n;K;01.09.2026;;İade\n;K;01.09.2026;;Kayıp\n7;;01.09.2026;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["durum"], x["teslim"]) for x in l], sorted(u["tur"] for u in uy)),
                             ([("1", "Teslim edildi", date(2026, 9, 3)), ("2", "Yolda", None), ("6", "Teslim edildi", None), ("", "İade", None), ("", "Kayıp", None)],
                              ["Gönderi iki kez yazılmış", "Satır okunamadı", "Satır okunamadı", "Teslim tarihi yok", "Teslim verilişten önce"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--gonderiler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "30.09.2026", "--hedef", "2", "--gecikme", "3", "--zamaninda", "95", "--hasar", "0,5", "--kayip", "0"]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--hedef", "-1"], ["--hasar", "101"])], [2] * 3)
            self.assertEqual(main.main(["--gonderiler", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Takip No;Taşıyıcı\n1;K\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
