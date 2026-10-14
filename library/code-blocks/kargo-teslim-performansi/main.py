"""
Kargo Teslim Performansı — Workers / Workless kod bloğu
E-ticaret Operasyon ve Kargo › E-ticaret Operasyon Uzmanı
Kargo firması bazında teslim süresi, zamanında teslim, hasarlı, kayıp ve iade gönderi oranlarını hesaplar.
  Teslim süresi        = teslim tarihi − kargoya veriliş tarihi (takvim günü)
  Zamanında teslim %   = teslim süresi --hedef (3) gün ve altı olan gönderi ÷ teslim edilen gönderi × 100
  Hasarlı / kayıp / iade oranı % = ilgili gönderi ÷ toplam gönderi × 100
  Gecikmiş açık gönderi = teslim edilmemiş, kayıp ya da iade olmamış ve kargoya verilişinden --gecikme (5) günden fazla geçmiş
  Uyarı                = zamanında teslim --zamaninda (%90) altında · hasar oranı --hasar (%1) üstünde · kayıp oranı --kayip (%0,5) üstünde
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"no": ("gonderi no", "takip no", "kargo takip no", "siparis no", "barkod"), "firma": ("kargo firmasi", "kargo", "firma", "tasiyici"),
         "verilis": ("kargoya verilis tarihi", "kargoya verilis", "gonderim tarihi", "cikis tarihi", "teslim alma tarihi"), "teslim": ("teslim tarihi", "teslimat tarihi", "alici teslim tarihi"),
         "durum": ("durum", "gonderi durumu", "son durum"), "il": ("il", "sehir", "varis ili", "alici ili")}
DURUM = {"teslim edildi": "Teslim edildi", "teslim": "Teslim edildi", "teslim edilmis": "Teslim edildi", "hasarli": "Hasarlı", "hasarli teslim": "Hasarlı", "kayip": "Kayıp",
         "iade": "İade", "iade edildi": "İade", "geri dondu": "İade", "yolda": "Yolda", "dagitimda": "Yolda", "transferde": "Yolda", "subede": "Yolda", "teslim edilemedi": "Yolda"}


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("firma", "verilis")):
        no, firma = metin(r["no"]), metin(r["firma"])
        if not firma:
            continue
        v, t = tarih(r["verilis"]), tarih(r["teslim"]) if metin(r["teslim"]) else None
        kd = katla(r["durum"])
        durum = DURUM.get(kd) if kd else ("Teslim edildi" if t else "Yolda")
        kim = f"{firma} · {no or r['_satir']}"
        if v is None or (metin(r["teslim"]) and t is None) or durum is None:
            uy.append(uyari("Yüksek", "Satır okunamadı", kim, f"Satır {r['_satir']}: kargoya veriliş tarihi, teslim tarihi ya da durum ('{metin(r['durum'])}') okunamadı. Satır alınmadı."))
            continue
        if t and t < v:
            uy.append(uyari("Yüksek", "Teslim verilişten önce", kim, f"Satır {r['_satir']}: teslim {t:%d.%m.%Y}, veriliş {v:%d.%m.%Y}. Satır alınmadı."))
            continue
        if no and katla(no) in gorulen:
            uy.append(uyari("Orta", "Gönderi iki kez yazılmış", kim, f"Satır {r['_satir']}: ikinci satır alınmadı."))
            continue
        if no:
            gorulen.add(katla(no))
        if durum in ("Teslim edildi", "Hasarlı") and t is None:
            uy.append(uyari("Bilgi", "Teslim tarihi yok", kim, "Teslim edilmiş görünüyor ama teslim tarihi yazılmamış; süre hesabına girmedi."))
        l.append({"no": no, "firma": firma, "verilis": v, "teslim": t, "durum": durum, "il": metin(r["il"])})
    return l, uy


def _medyan(l: list[int]) -> Decimal | None:
    if not l:
        return None
    s, n = sorted(l), len(l)
    return Decimal(s[n // 2]) if n % 2 else Decimal(s[n // 2 - 1] + s[n // 2]) / 2


def gosterge(l: list[dict], hedef: int) -> dict:
    sure = [(g["teslim"] - g["verilis"]).days for g in l if g["teslim"] and g["durum"] in ("Teslim edildi", "Hasarlı")]
    n = len(l)
    say = lambda d: sum(1 for g in l if g["durum"] == d)   # noqa: E731
    oran = lambda x: yuvarla(Decimal(x) * 100 / n) if n else None   # noqa: E731
    return {"gonderi": n, "teslim": len(sure), "ort": yuvarla(Decimal(sum(sure)) / len(sure)) if sure else None, "medyan": _medyan(sure), "en_uzun": max(sure) if sure else None,
            "zamaninda": yuvarla(Decimal(sum(1 for s in sure if s <= hedef)) * 100 / len(sure)) if sure else None, "hasar": say("Hasarlı"), "hasar_oran": oran(say("Hasarlı")),
            "kayip": say("Kayıp"), "kayip_oran": oran(say("Kayıp")), "iade": say("İade"), "iade_oran": oran(say("İade")), "acik": say("Yolda")}


def analiz(gonderiler: list[dict], gun: date, hedef: int = 3, gecikme: int = 5, zamaninda: Decimal = Decimal(90), hasar: Decimal = Decimal(1), kayip: Decimal = Decimal("0.5")) -> dict:
    if not gonderiler:
        raise ValueError("Gönderi kaydı yok.")
    uy, firmalar, iller = [], {}, {}
    for g in gonderiler:
        firmalar.setdefault(katla(g["firma"]), []).append(g)
        if g["il"]:
            iller.setdefault((katla(g["firma"]), katla(g["il"])), []).append(g)
    sonuc = []
    for k in sorted(firmalar):
        l = firmalar[k]
        f = {"firma": l[0]["firma"], **gosterge(l, hedef)}
        if f["zamaninda"] is not None and f["zamaninda"] < zamaninda:
            uy.append(uyari("Orta", "Zamanında teslim düşük", f["firma"], f"{f['teslim']} teslimin %{f['zamaninda']}'i {hedef} gün içinde (eşik %{zamaninda}); ortalama {f['ort']} gün."))
        if f["hasar_oran"] > hasar:
            uy.append(uyari("Orta", "Hasar oranı yüksek", f["firma"], f"{f['gonderi']} gönderide {f['hasar']} hasarlı (%{f['hasar_oran']}; eşik %{hasar})."))
        if f["kayip_oran"] > kayip:
            uy.append(uyari("Yüksek", "Kayıp oranı yüksek", f["firma"], f"{f['gonderi']} gönderide {f['kayip']} kayıp (%{f['kayip_oran']}; eşik %{kayip})."))
        sonuc.append(f)
    acik = []
    for g in gonderiler:
        if g["durum"] == "Yolda":
            yas = (gun - g["verilis"]).days
            if yas > gecikme:
                acik.append({**g, "yas": yas})
                uy.append(uyari("Orta", "Gecikmiş gönderi", f"{g['firma']} · {g['no']}", f"{g['verilis']:%d.%m.%Y} tarihinde kargoya verildi, {yas} gündür teslim edilmedi" + (f" ({g['il']})." if g["il"] else ".")))
    acik.sort(key=lambda g: (-g["yas"], katla(g["firma"]), g["no"]))
    il = [{"firma": l[0]["firma"], "il": l[0]["il"], **gosterge(l, hedef)} for _, l in sorted(iller.items())]
    return {"firmalar": sonuc, "toplam": gosterge(gonderiler, hedef), "iller": il, "acik": acik, "uyarilar": uy}


BASLIK = ["Gönderi", "Teslim Edilen", "Ortalama Süre (gün)", "Medyan Süre", "En Uzun", "Zamanında %", "Hasarlı", "Hasar %", "Kayıp", "Kayıp %", "İade", "İade %", "Açık"]


def _deger(f: dict) -> list:
    return [f["gonderi"], f["teslim"], f["ort"], f["medyan"], f["en_uzun"], f["zamaninda"], f["hasar"], f["hasar_oran"], f["kayip"], f["kayip_oran"], f["iade"], f["iade_oran"], f["acik"]]


def rapor_yaz(cikti: Path, r: dict, gun: date, hedef: int, zamaninda: Decimal, hasar: Decimal, kayip: Decimal) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Firmalar"
    _baslik(ws, ["Kargo Firması"] + BASLIK, (20, 9, 9, 10, 9, 8, 10, 8, 8, 7, 8, 7, 8, 7), "B2")

    def renk(f):
        return {7: None if f["zamaninda"] is None else KIRMIZI if f["zamaninda"] < zamaninda else YESIL, 9: KIRMIZI if f["hasar_oran"] and f["hasar_oran"] > hasar else None,
                11: KIRMIZI if f["kayip_oran"] and f["kayip_oran"] > kayip else None}
    for f in r["firmalar"]:
        _satir(ws, [f["firma"]] + _deger(f), para_s=(4, 5, 7, 9, 11, 13), dolgu=renk(f))
    k = _satir(ws, ["TOPLAM"] + _deger(r["toplam"]), para_s=(4, 5, 7, 9, 11, 13))
    for h in ws[k]:
        h.font = Font(bold=True)
    ws.append([])
    ws.append([f"Süreler takvim günüdür (hafta sonu ve tatil dahil). Zamanında: {hedef} gün ve altı. Rapor tarihi {gun:%d.%m.%Y}."])
    if r["iller"]:
        ws = wb.create_sheet("Firma ve İl")
        _baslik(ws, ["Kargo Firması", "İl"] + BASLIK, (20, 14, 9, 9, 10, 9, 8, 10, 8, 8, 7, 8, 7, 8, 7), "C2")
        for f in r["iller"]:
            _satir(ws, [f["firma"], f["il"]] + _deger(f), para_s=(5, 6, 8, 10, 12, 14), dolgu={k + 1: v for k, v in renk(f).items()})
    ws = wb.create_sheet("Gecikmiş Gönderiler")
    _baslik(ws, ["Gönderi No", "Kargo Firması", "İl", "Kargoya Veriliş", "Geçen Gün"], (18, 20, 14, 12, 9))
    for g in r["acik"]:
        _satir(ws, [g["no"], g["firma"], g["il"], g["verilis"], g["yas"]], tarih_s=(4,))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, gun: date, hedef: int = 3, gecikme: int = 5, zamaninda: Decimal = Decimal(90), hasar: Decimal = Decimal(1), kayip: Decimal = Decimal("0.5")) -> dict:
    gonderiler, uy = oku(yol)
    r = analiz(gonderiler, gun, hedef, gecikme, zamaninda, hasar, kayip)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, gun, hedef, zamaninda, hasar, kayip)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Kargo firması bazında teslim süresi, hasarlı ve kayıp gönderi oranlarını hesaplar.")
    p.add_argument("--gonderiler", type=Path, help="Gönderi No, Kargo Firması, Kargoya Veriliş Tarihi, Teslim Tarihi, Durum, İl")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY; açık gönderilerin yaşı için (varsayılan bugün; örnek veride 30.09.2026)")
    p.add_argument("--hedef", type=int, default=3, help="Zamanında teslim sayılacak en uzun süre, gün (varsayılan 3)")
    p.add_argument("--gecikme", type=int, default=5, help="Açık gönderi bu kadar günü geçince gecikmiş sayılır (varsayılan 5)")
    p.add_argument("--zamaninda", default="90", help="Zamanında teslim oranı bu yüzdenin altındaysa uyarılır (varsayılan 90)")
    p.add_argument("--hasar", default="1", help="Hasar oranı bu yüzdenin üstündeyse uyarılır (varsayılan 1)")
    p.add_argument("--kayip", default="0,5", help="Kayıp oranı bu yüzdenin üstündeyse uyarılır (varsayılan 0,5)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kargo_teslim_performansi.xlsx")
    a = p.parse_args(argv)
    yol = a.gonderiler or ORNEK / "kargo_gonderileri.csv"
    gun = tarih(a.tarih) if a.tarih else (date(2026, 9, 30) if a.gonderiler is None else date.today())
    oranlar = [para(x) for x in (a.zamaninda, a.hasar, a.kayip)]
    if gun is None or a.hedef < 0 or a.gecikme < 0 or any(x is None or not 0 <= x <= 100 for x in oranlar):
        print("[X] --tarih GG.AA.YYYY, gün sayıları sıfır ya da üstü, oranlar 0-100 arası olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, gun, a.hedef, a.gecikme, *oranlar)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {len(r['firmalar'])} firma · {t['gonderi']} gönderi · ortalama {t['ort']} gün · zamanında %{t['zamaninda']} · hasar %{t['hasar_oran']} · kayıp %{t['kayip_oran']} · gecikmiş açık {len(r['acik'])}")
    for f in r["firmalar"]:
        print(f"     {f['firma']}: {f['gonderi']} gönderi · ortalama {f['ort']} gün · zamanında %{f['zamaninda']} · hasar %{f['hasar_oran']} · kayıp %{f['kayip_oran']}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek" or (u["onem"] == "Orta" and u["tur"] != "Gecikmiş gönderi"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
