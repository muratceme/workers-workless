"""
Mizandan Mali Tablo Hazırlama — Workers / Workless kod bloğu
Muhasebe › Muhasebe Müdürü

Tekdüzen Hesap Planı'na göre tutulan mizandan, Muhasebe Sistemi Uygulama Genel Tebliği (MSUGT) biçiminde
ayrıntılı bilanço ve gelir tablosu hazırlar; önceki dönem mizanı verilirse karşılaştırmalı sütun ekler.
  - Hesaplar ana hesap (3 hane) düzeyinde, grup başlıkları (A. Hazır Değerler, B. Menkul Kıymetler …) ve
    toplamlarla gösterilir; düzenleyici (-) hesaplar eksi tutarla yer alır.
  - Ters bakiyeli alt hesaplar bilançoda doğru tarafa virmanlanır (müşteri avansı → 340, satıcıya avans → 159,
    kredili mevduat → 300, personel avansı → 196); yapılan virmanlar ayrı sayfada listelenir.
  - Gelir tablosu hesapları açıksa dönem kârı hesaplanıp özkaynaklarda 590/591 olarak gösterilir.
  - Kontroller: aktif = pasif, 7'li hesaplarda kalan bakiye, kapanış hesapları (690/692/697/698).
İnternete bağlanmaz.

Kullanım:
    python main.py                                         # örnek mizanlarla dener
    python main.py --mizan mizan_2026.xlsx --onceki mizan_2025.xlsx --unvan "Örnek A.Ş."
    python main.py --mizan mizan_2026.xlsx --virman-yok    # ters bakiyeleri olduğu yerde bırak
"""
from __future__ import annotations

import argparse
import sys
from collections import OrderedDict
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import mizan_cekirdek as mc

BURASI = Path(__file__).resolve().parent
SIFIR = Decimal(0)

# MSUGT ayrıntılı bilanço düzeni: (bölüm, [(grup harfi, grup adı, hesap grubu önekleri)]); pasif III-V diye sürer
AKTIF = [
    ("I. DÖNEN VARLIKLAR", [("A", "Hazır Değerler", ("10",)), ("B", "Menkul Kıymetler", ("11",)),
                             ("C", "Ticari Alacaklar", ("12",)), ("D", "Diğer Alacaklar", ("13",)),
                             ("E", "Stoklar", ("15",)), ("F", "Yıllara Yaygın İnşaat ve Onarım Maliyetleri", ("17",)),
                             ("G", "Gelecek Aylara Ait Giderler ve Gelir Tahakkukları", ("18",)),
                             ("H", "Diğer Dönen Varlıklar", ("19",))]),
    ("II. DURAN VARLIKLAR", [("A", "Ticari Alacaklar", ("22",)), ("B", "Diğer Alacaklar", ("23",)),
                              ("C", "Mali Duran Varlıklar", ("24",)), ("D", "Maddi Duran Varlıklar", ("25",)),
                              ("E", "Maddi Olmayan Duran Varlıklar", ("26",)), ("F", "Özel Tükenmeye Tabi Varlıklar", ("27",)),
                              ("G", "Gelecek Yıllara Ait Giderler ve Gelir Tahakkukları", ("28",)),
                              ("H", "Diğer Duran Varlıklar", ("29",))]),
]
PASIF = [
    ("III. KISA VADELİ YABANCI KAYNAKLAR", [("A", "Mali Borçlar", ("30",)), ("B", "Ticari Borçlar", ("32",)),
                                          ("C", "Diğer Borçlar", ("33",)), ("D", "Alınan Avanslar", ("34",)),
                                          ("E", "Yıllara Yaygın İnşaat ve Onarım Hakedişleri", ("35",)),
                                          ("F", "Ödenecek Vergi ve Diğer Yükümlülükler", ("36",)),
                                          ("G", "Borç ve Gider Karşılıkları", ("37",)),
                                          ("H", "Gelecek Aylara Ait Gelirler ve Gider Tahakkukları", ("38",)),
                                          ("I", "Diğer Kısa Vadeli Yabancı Kaynaklar", ("39",))]),
    ("IV. UZUN VADELİ YABANCI KAYNAKLAR", [("A", "Mali Borçlar", ("40",)), ("B", "Ticari Borçlar", ("42",)),
                                           ("C", "Diğer Borçlar", ("43",)), ("D", "Alınan Avanslar", ("44",)),
                                           ("E", "Borç ve Gider Karşılıkları", ("47",)),
                                           ("F", "Gelecek Yıllara Ait Gelirler ve Gider Tahakkukları", ("48",)),
                                           ("G", "Diğer Uzun Vadeli Yabancı Kaynaklar", ("49",))]),
    ("V. ÖZKAYNAKLAR", [("A", "Ödenmiş Sermaye", ("50",)), ("B", "Sermaye Yedekleri", ("52",)),
                          ("C", "Kâr Yedekleri", ("54",)), ("D", "Geçmiş Yıllar Kârları", ("57",)),
                          ("E", "Geçmiş Yıllar Zararları (-)", ("58",)), ("F", "Dönem Net Kârı (Zararı)", ("59",))]),
]
GELIR = [  # (harf, ad, hesap önekleri; önek None ise ara toplam satırı)
    ("A", "Brüt Satışlar", ("60",)),
    ("B", "Satış İndirimleri (-)", ("61",)),
    ("C", "NET SATIŞLAR", None),
    ("D", "Satışların Maliyeti (-)", ("62",)),
    ("", "BRÜT SATIŞ KÂRI VEYA ZARARI", None),
    ("E", "Faaliyet Giderleri (-)", ("63",)),
    ("", "FAALİYET KÂRI VEYA ZARARI", None),
    ("F", "Diğer Faaliyetlerden Olağan Gelir ve Kârlar", ("64",)),
    ("G", "Diğer Faaliyetlerden Olağan Gider ve Zararlar (-)", ("65",)),
    ("H", "Finansman Giderleri (-)", ("66",)),
    ("", "OLAĞAN KÂR VEYA ZARAR", None),
    ("I", "Olağandışı Gelir ve Kârlar", ("67",)),
    ("J", "Olağandışı Gider ve Zararlar (-)", ("68",)),
    ("", "DÖNEM KÂRI VEYA ZARARI", None),
    ("K", "Dönem Kârı Vergi ve Diğer Yasal Yükümlülük Karşılıkları (-)", ("691",)),
    ("", "DÖNEM NET KÂRI VEYA ZARARI", None),
]

# Ters bakiyeli alt hesapların bilançodaki karşı hesabı: (ana hesap önekleri, beklenen doğa, hedef hesap, hedef adı)
VIRMAN = [
    (("102",), "B", "300", "Banka Kredileri (kredili mevduat)"),
    (("120", "121", "127"), "B", "340", "Alınan Sipariş Avansları"),
    (("220", "221"), "B", "440", "Alınan Sipariş Avansları"),
    (("320", "321", "329"), "A", "159", "Verilen Sipariş Avansları"),
    (("420", "421"), "A", "259", "Verilen Avanslar"),
    (("335",), "A", "196", "Personel Avansları"),
]


def virmanla(h: dict[str, dict]) -> tuple[dict[str, dict], list[dict]]:
    """En alt kırılımda ters bakiye veren cari/banka hesaplarını karşı tarafa taşır. Ana hesap bakiyeleri döner."""
    ana = mc.ana_hesaplar(h)
    kodlar = set(h)
    cocuklu = {mc.ust_kod(k, kodlar) for k in kodlar} - {None}
    virmanlar = []
    for k, x in sorted(h.items()):
        if k == k[:3] or k in cocuklu:
            continue
        for onekler, beklenen, hedef, hedef_ad in VIRMAN:
            if not k.startswith(onekler):
                continue
            ters = (beklenen == "B" and x["net"] < 0) or (beklenen == "A" and x["net"] > 0)
            if ters:
                ana[k[:3]]["net"] -= x["net"]
                hd = ana.setdefault(hedef, {"kod": hedef, "ad": hedef_ad, "net": SIFIR, "borc": None, "alacak": None})
                hd["net"] += x["net"]
                hd["ad"] = hd["ad"] or hedef_ad
                virmanlar.append({"kod": k, "ad": x["ad"], "tutar": x["net"], "hedef": hedef, "hedef_ad": hedef_ad})
    return ana, virmanlar


def _gt(m: dict[str, dict], onekler) -> Decimal:
    """Gelir tablosu tutarı: alacak − borç (gelir +, gider −)."""
    return -sum((v["net"] for k, v in m.items() if k.startswith(onekler)), SIFIR)


def gelir_tablosu(m: dict[str, dict]) -> tuple[list[tuple], Decimal]:
    satirlar, ara = [], SIFIR
    for harf, ad, onek in GELIR:
        if onek is None:
            satirlar.append((harf, ad, None, ara, "ara"))
            continue
        hesaplar = [(k, v["ad"], -v["net"]) for k, v in sorted(m.items()) if k.startswith(onek) and v["net"]]
        tutar = _gt(m, onek)
        ara += tutar
        satirlar.append((harf, ad, hesaplar, tutar, "grup"))
    return satirlar, ara


def tablolar(m: dict[str, dict]) -> tuple[dict, list[str]]:
    uyarilar = []
    acik_6 = any(v["net"] for k, v in m.items() if k[0] == "6" and not k.startswith(("690", "692")))
    gt, donem_net = gelir_tablosu(m)
    m = {k: dict(v) for k, v in m.items()}
    if acik_6:
        if any(m.get(k, {}).get("net") for k in ("590", "591")):
            uyarilar.append("Hem gelir tablosu hesapları açık hem 590/591'de bakiye var: dönem kârı çift sayılabilir")
        kod = "590" if donem_net >= 0 else "591"
        m[kod] = {"kod": kod, "ad": "Dönem Net Kârı" if donem_net >= 0 else "Dönem Net Zararı (-)",
                  "net": (m.get(kod, {}).get("net") or SIFIR) - donem_net, "borc": None, "alacak": None}
    else:
        uyarilar.append("Gelir tablosu hesapları kapalı (kapanış sonrası mizan): gelir tablosu boş, dönem kârı 590/591'den alındı")
        donem_net = -sum((v["net"] for k, v in m.items() if k.startswith("59")), SIFIR)
    for k in ("690", "692", "697", "698"):
        if m.get(k, {}).get("net"):
            uyarilar.append(f"{k} hesabında bakiye var: kapanış kayıtları tamamlanmamış olabilir")
    yedi = sum((v["net"] for k, v in m.items() if k[0] == "7"), SIFIR)
    if abs(yedi) > 1:
        uyarilar.append(f"7'li maliyet hesaplarında {mc.tl(yedi)} bakiye var (yansıtma yapılmamış); bilanço denkliği bozulabilir")

    def bolum(yapi, pasif: bool):
        sonuc = []
        for baslik, gruplar in yapi:
            gsat, btop = [], SIFIR
            for harf, ad, onek in gruplar:
                hesaplar = [(k, v["ad"], -v["net"] if pasif else v["net"]) for k, v in sorted(m.items())
                            if k.startswith(onek) and len(k) == 3 and v["net"]]
                top = sum((t for _, _, t in hesaplar), SIFIR)
                btop += top
                gsat.append((harf, ad, hesaplar, top))
            sonuc.append((baslik, gsat, btop))
        return sonuc

    aktif, pasif = bolum(AKTIF, False), bolum(PASIF, True)
    a_top = sum((b[2] for b in aktif), SIFIR)
    p_top = sum((b[2] for b in pasif), SIFIR)
    if abs(a_top - p_top) > 1:
        uyarilar.append(f"Bilanço denk değil: aktif {mc.tl(a_top)} ≠ pasif {mc.tl(p_top)} (fark {mc.tl(a_top - p_top)})")
    return {"aktif": aktif, "pasif": pasif, "aktif_top": a_top, "pasif_top": p_top, "gelir": gt,
            "donem_net": donem_net, "acik_6": acik_6}, uyarilar


def calistir(mizan: Path, cikti: Path, onceki: Path | None = None, unvan: str = "", virman: bool = True) -> dict:
    donemler = OrderedDict()
    for etiket, yol in (("Cari Dönem", mizan), ("Önceki Dönem", onceki)):
        if not yol:
            continue
        h = mc.mizan_oku(yol)
        ana, virmanlar = virmanla(h) if virman else (mc.ana_hesaplar(h), [])
        t, u = tablolar(ana)
        donemler[etiket] = {"yol": yol, "t": t, "uyarilar": u, "virmanlar": virmanlar}
    _rapor(donemler, cikti, unvan)
    return donemler


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

PARA = '#,##0.00;(#,##0.00);"-"'
KALIN = Font(bold=True)
BOLUM_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BOLUM_YAZI = Font(bold=True, color="FFFFFF")
GRUP_DOLGU = PatternFill("solid", fgColor="EEF1F6")
TOPLAM_DOLGU = PatternFill("solid", fgColor="DDE5F0")
UST_CIZGI = Border(top=Side(style="thin"))


def _hesap_esle(donemler, secici):
    """Dönemler arası aynı satırları hizalamak için: tüm dönemlerdeki hesap kodlarının birleşimi."""
    kodlar = OrderedDict()
    for d in donemler.values():
        for k, ad, _ in secici(d["t"]):
            kodlar.setdefault(k, ad)
    return kodlar


def _bilanco_sayfasi(ws, baslik, donemler, anahtar, toplam_anahtar, toplam_ad):
    etiketler = list(donemler)
    ws.append([baslik] + [""] + etiketler)
    for c in ws[ws.max_row]:
        c.fill, c.font = BOLUM_DOLGU, BOLUM_YAZI
    ilk = next(iter(donemler.values()))["t"][anahtar]
    for bi, (bolum_ad, gruplar, _) in enumerate(ilk):
        ws.append([bolum_ad, ""] + [float(d["t"][anahtar][bi][2]) for d in donemler.values()])
        for c in ws[ws.max_row]:
            c.font, c.fill = KALIN, TOPLAM_DOLGU
        for gi, (harf, ad, _, _) in enumerate(gruplar):
            toplamlar = [d["t"][anahtar][bi][1][gi][3] for d in donemler.values()]
            hesaplar = OrderedDict()
            for d in donemler.values():
                for k, had, _ in d["t"][anahtar][bi][1][gi][2]:
                    hesaplar.setdefault(k, had)
            if not hesaplar and not any(toplamlar):
                continue
            ws.append([f"   {harf}. {ad}", ""] + [float(t) for t in toplamlar])
            for c in ws[ws.max_row]:
                c.font, c.fill = KALIN, GRUP_DOLGU
            for k, had in hesaplar.items():
                tutarlar = []
                for d in donemler.values():
                    eslesen = [t for kk, _, t in d["t"][anahtar][bi][1][gi][2] if kk == k]
                    tutarlar.append(float(eslesen[0]) if eslesen else None)
                ws.append([f"        {k} {had}", k] + tutarlar)
    ws.append([toplam_ad, ""] + [float(d["t"][toplam_anahtar]) for d in donemler.values()])
    for c in ws[ws.max_row]:
        c.font, c.fill, c.border = Font(bold=True, color="FFFFFF"), BOLUM_DOLGU, UST_CIZGI
    for r in ws.iter_rows(min_row=2):
        for c in r[2:]:
            c.number_format = PARA
    ws.column_dimensions["A"].width = 62
    ws.column_dimensions["B"].width = 7
    for j in range(len(etiketler)):
        ws.column_dimensions[get_column_letter(3 + j)].width = 20
    ws.freeze_panes = "C2"


def _rapor(donemler, cikti, unvan):
    wb = Workbook()
    ilk_ad = next(iter(donemler.values()))["yol"].name
    o = wb.active
    o.title = "Bilgi"
    o.append([f"{unvan or 'Şirket'} — Mali Tablolar"])
    o["A1"].font = Font(bold=True, size=13)
    o.append(["Kaynak", ", ".join(f"{e}: {d['yol'].name}" for e, d in donemler.items())])
    o.append(["Biçim", "Muhasebe Sistemi Uygulama Genel Tebliği (Tekdüzen Hesap Planı) ayrıntılı bilanço ve gelir tablosu"])
    for e, d in donemler.items():
        t = d["t"]
        o.append([f"{e}: aktif / pasif", f"{mc.tl(t['aktif_top'])} / {mc.tl(t['pasif_top'])}"])
        o.append([f"{e}: dönem net kârı", mc.tl(t["donem_net"])])
        for u in d["uyarilar"]:
            o.append([f"{e}: uyarı", u])
    o.append(["Not", "Tutarlar mizandaki bakiyelerdir; dönem sonu değerleme, karşılık, amortisman ve kapanış kayıtları "
                      "mizana işlenmiş olmalıdır. Bu tablolar VUK/MSUGT biçimindedir; TFRS veya BOBİ FRS raporlaması değildir. "
                      "Beyanname ve yasal defterler için mali müşavirinizin kontrolünden geçirin."])
    o.column_dimensions["A"].width = 30
    o.column_dimensions["B"].width = 120
    for r in o.iter_rows():
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    _bilanco_sayfasi(wb.create_sheet("Bilanço Aktif"), "AKTİF (VARLIKLAR)", donemler, "aktif", "aktif_top", "AKTİF TOPLAMI")
    _bilanco_sayfasi(wb.create_sheet("Bilanço Pasif"), "PASİF (KAYNAKLAR)", donemler, "pasif", "pasif_top", "PASİF TOPLAMI")

    g = wb.create_sheet("Gelir Tablosu")
    g.append(["GELİR TABLOSU", ""] + list(donemler))
    for c in g[1]:
        c.fill, c.font = BOLUM_DOLGU, BOLUM_YAZI
    ilk = next(iter(donemler.values()))["t"]["gelir"]
    for i, (harf, ad, _, _, tur) in enumerate(ilk):
        tutarlar = [float(d["t"]["gelir"][i][3]) for d in donemler.values()]
        if tur == "ara":
            g.append([f"{harf}. {ad}" if harf else ad, ""] + tutarlar)
            for c in g[g.max_row]:
                c.font, c.fill, c.border = KALIN, TOPLAM_DOLGU, UST_CIZGI
            continue
        g.append([f"{harf}. {ad}", ""] + tutarlar)
        for c in g[g.max_row]:
            c.font, c.fill = KALIN, GRUP_DOLGU
        hesaplar = OrderedDict()
        for d in donemler.values():
            for k, had, _ in d["t"]["gelir"][i][2]:
                hesaplar.setdefault(k, had)
        for k, had in hesaplar.items():
            satir = []
            for d in donemler.values():
                es = [t for kk, _, t in d["t"]["gelir"][i][2] if kk == k]
                satir.append(float(es[0]) if es else None)
            g.append([f"     {k} {had}", k] + satir)
    for r in g.iter_rows(min_row=2):
        for c in r[2:]:
            c.number_format = PARA
    g.column_dimensions["A"].width = 62
    g.column_dimensions["B"].width = 7
    for j in range(len(donemler)):
        g.column_dimensions[get_column_letter(3 + j)].width = 20
    g.freeze_panes = "C2"

    v = wb.create_sheet("Virmanlar")
    v.append(["Dönem", "Alt Hesap", "Hesap Adı", "Ters Bakiye (B+ / A−)", "Bilançoda Gösterildiği Hesap"])
    for c in v[1]:
        c.fill, c.font = BOLUM_DOLGU, BOLUM_YAZI
    for e, d in donemler.items():
        for x in d["virmanlar"]:
            v.append([e, x["kod"], x["ad"], float(x["tutar"]), f"{x['hedef']} {x['hedef_ad']}"])
            v.cell(v.max_row, 4).number_format = PARA
    if v.max_row == 1:
        v.append(["", "Virman yapılmadı"])
    for j, w in enumerate((14, 14, 32, 20, 40), 1):
        v.column_dimensions[get_column_letter(j)].width = w
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ornek = BURASI / "ornek_veri"
    ap = argparse.ArgumentParser(description="Mizandan Tekdüzen (MSUGT) biçiminde ayrıntılı bilanço ve gelir tablosu hazırlar.")
    ap.add_argument("--mizan", type=Path, default=ornek / "mizan_2026.csv", help="Cari dönem mizanı (.xlsx/.csv)")
    ap.add_argument("--onceki", type=Path, help="Önceki dönem mizanı (karşılaştırmalı sütun için)")
    ap.add_argument("--unvan", default="", help="Şirket unvanı (rapor başlığı)")
    ap.add_argument("--virman-yok", action="store_true", help="Ters bakiyeli alt hesapları karşı tarafa taşıma")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "mali_tablolar.xlsx")
    a = ap.parse_args(argv)
    if a.mizan == ornek / "mizan_2026.csv" and not a.onceki:
        a.onceki = ornek / "mizan_2025.csv"
        a.unvan = a.unvan or "Örnek Sanayi A.Ş. (kurgusal)"
    s = calistir(a.mizan, a.cikti, a.onceki, a.unvan, not a.virman_yok)
    for e, d in s.items():
        t = d["t"]
        print(f"[OK] {e}: aktif {mc.tl(t['aktif_top'])} · pasif {mc.tl(t['pasif_top'])} · dönem net kârı {mc.tl(t['donem_net'])}"
              f" · virman {len(d['virmanlar'])}")
        for u in d["uyarilar"]:
            print(f"[!] {e}: {u}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
