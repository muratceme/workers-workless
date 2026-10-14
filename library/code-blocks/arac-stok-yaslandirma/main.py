"""
Araç Stok Yaşlandırma — Workers / Workless kod bloğu
Satış › Satış Müdürü
Stoktaki araçları stoğa giriş tarihine göre yaşlandırır; model, renk ve yaş bandı bazında adet, değer ve finansman maliyetini hesaplar.
  Stok günü          = rapor tarihi − stoğa giriş tarihi (satılmış / teslim edilmiş araçlar alınmaz)
  Yaş bandı          = 0–30 · 31–60 · 61–90 · 91–180 · 180+ gün
  Finansman maliyeti = alış maliyeti × yıllık oran ÷ 100 × max(stok günü − faizsiz gün, 0) ÷ 365 (basit faiz)
                       oran: floor plan / kredi ile alınanlara --faiz, özkaynakla alınanlara --firsat-faiz (fırsat maliyeti)
  Marj tüketimi %    = finansman maliyeti ÷ (liste fiyatı − alış maliyeti) × 100
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
SUTUN = {"sasi": ("sasi no", "sasi", "vin", "sase no"), "marka": ("marka",), "model": ("model", "model adi"), "versiyon": ("versiyon", "donanim", "paket"), "renk": ("renk", "dis renk"),
         "yil": ("model yili", "yil"), "giris": ("stoga giris tarihi", "stok giris tarihi", "giris tarihi", "fatura tarihi", "teslim alma tarihi"), "alis": ("alis maliyeti", "alis fiyati", "maliyet", "alis"),
         "liste": ("liste fiyati", "satis fiyati", "tavsiye fiyat"), "durum": ("durum", "stok durumu"), "finansman": ("finansman", "finansman tipi", "odeme sekli"),
         "faizsiz": ("faizsiz gun", "faizsiz sure", "vade gun", "bedelsiz gun")}
BANT = (("0–30 gün", 30), ("31–60 gün", 60), ("61–90 gün", 90), ("91–180 gün", 180), ("180+ gün", None))
SIFIR = Decimal(0)


def _durum(x) -> str | None:
    k = katla(x)
    if not k or k in ("stokta", "stok", "musait", "satista", "yolda"):
        return "Stokta"
    if k in ("rezerve", "rezerv", "kaparo", "kaporali", "siparisli", "bagli"):
        return "Rezerve"
    if k in ("satildi", "teslim edildi", "faturalandi", "satis", "teslim"):
        return "Satıldı"
    return None


def _ozkaynak(x) -> bool:
    k = katla(x)
    return k in ("ozkaynak", "nakit", "pesin", "kendi kaynak", "ozsermaye")


def bant(gun: int) -> str:
    return next(ad for ad, ust in BANT if ust is None or gun <= ust)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("sasi", "giris", "alis")):
        sasi = "".join(metin(r["sasi"]).upper().split())
        if not sasi:
            continue
        g, alis, durum = tarih(r["giris"]), para(r["alis"]), _durum(r["durum"])
        liste = para(r["liste"]) if metin(r["liste"]) else None
        faizsiz = para(r["faizsiz"]) if metin(r["faizsiz"]) else SIFIR
        yil = para(r["yil"]) if metin(r["yil"]) else None
        if g is None or alis is None or durum is None or faizsiz is None or (metin(r["liste"]) and liste is None) or (metin(r["yil"]) and yil is None):
            uy.append(uyari("Yüksek", "Satır okunamadı", sasi, f"Satır {r['_satir']}: giriş tarihi, alış maliyeti, durum ('{metin(r['durum'])}'), liste fiyatı, model yılı ya da faizsiz gün okunamadı."))
            continue
        if sasi in gorulen:
            uy.append(uyari("Yüksek", "Şasi iki kez yazılmış", sasi, f"Satır {r['_satir']}: ikinci satır alınmadı."))
            continue
        gorulen.add(sasi)
        if len(sasi) != 17 or any(c in "IOQ" for c in sasi):
            uy.append(uyari("Bilgi", "Şasi numarası biçimi", sasi, "Şasi numarası 17 karakter olmalı ve I, O, Q harflerini içermemeli; kaydı kontrol edin."))
        if durum == "Satıldı":
            continue
        l.append({"sasi": sasi, "marka": metin(r["marka"]), "model": metin(r["model"]) or "Belirtilmemiş", "versiyon": metin(r["versiyon"]), "renk": metin(r["renk"]) or "Belirtilmemiş",
                  "yil": int(yil) if yil is not None else None, "giris": g, "alis": alis, "liste": liste, "durum": durum, "finansman": metin(r["finansman"]),
                  "ozkaynak": _ozkaynak(r["finansman"]), "faizsiz": int(faizsiz)})
    return l, uy


def analiz(araclar: list[dict], bugun: date, faiz: Decimal | None = None, firsat: Decimal | None = None, marj_esik: Decimal = Decimal(50)) -> dict:
    if not araclar:
        raise ValueError("Stokta araç yok.")
    uy = []
    for a in araclar:
        kim = f"{a['sasi']} · {a['model']} {a['renk']}".strip()
        a["gun"] = (bugun - a["giris"]).days
        if a["gun"] < 0:
            uy.append(uyari("Orta", "Giriş tarihi ileride", kim, f"Stoğa giriş {a['giris']:%d.%m.%Y}; stok günü 0 sayıldı."))
            a["gun"] = 0
        a["bant"] = bant(a["gun"])
        oran = firsat if a["ozkaynak"] else faiz
        a["oran"] = oran
        a["faizli_gun"] = max(a["gun"] - a["faizsiz"], 0)
        a["maliyet"] = yuvarla(a["alis"] * oran / 100 * a["faizli_gun"] / 365) if oran is not None else None
        a["marj"] = a["liste"] - a["alis"] if a["liste"] is not None else None
        a["marj_tuketim"] = yuvarla(a["maliyet"] * 100 / a["marj"]) if a["maliyet"] is not None and a["marj"] and a["marj"] > 0 else None
        if a["gun"] > 180:
            uy.append(uyari("Yüksek", "180 günü geçen stok", kim, f"{a['gun']} gündür stokta" + (f"; finansman maliyeti {a['maliyet']:,.2f}" if a["maliyet"] else "") + ". Kampanya, transfer ya da fiyat kararı gerekli."))
        elif a["gun"] > 90:
            uy.append(uyari("Orta", "90 günü geçen stok", kim, f"{a['gun']} gündür stokta."))
        if a["oran"] and not a["ozkaynak"] and a["faizsiz"] and 0 <= a["faizsiz"] - a["gun"] <= 7:
            uy.append(uyari("Bilgi", "Faizsiz süre bitiyor", kim, f"Faizsiz süre {a['faizsiz'] - a['gun']} gün içinde bitiyor ({a['giris'] + timedelta(days=a['faizsiz']):%d.%m.%Y})."))
        if a["marj_tuketim"] is not None and a["marj_tuketim"] >= marj_esik:
            uy.append(uyari("Orta", "Finansman marjı eritiyor", kim, f"Finansman maliyeti {a['maliyet']:,.2f}, brüt marjın (liste − alış {a['marj']:,.2f}) %{a['marj_tuketim']} kadarı."))
        elif a["marj"] is not None and a["marj"] <= 0:
            uy.append(uyari("Bilgi", "Liste fiyatı maliyetin altında", kim, f"Liste {a['liste']:,.2f}, alış {a['alis']:,.2f}."))
        if a["yil"] and a["yil"] < bugun.year:
            uy.append(uyari("Bilgi", "Önceki model yılı", kim, f"Model yılı {a['yil']}; yıl dönümünde ikinci el değer ve kampanya etkisini değerlendirin."))
    if faiz is None:
        uy.append(uyari("Bilgi", "Faiz oranı verilmedi", "--faiz", "Floor plan / kredili araçlarda finansman maliyeti hesaplanmadı."))

    def ozet(fn, sira=None):
        g = {}
        for a in araclar:
            g.setdefault(fn(a), []).append(a)
        satir = []
        for ad, l in g.items():
            mal = [a["maliyet"] for a in l if a["maliyet"] is not None]
            satir.append({"ad": ad, "adet": len(l), "rezerve": sum(1 for a in l if a["durum"] == "Rezerve"), "alis": sum((a["alis"] for a in l), SIFIR),
                          "ort_gun": yuvarla(Decimal(sum(a["gun"] for a in l)) / len(l), "0.1"), "en_eski": max(a["gun"] for a in l), "maliyet": sum(mal, SIFIR) if mal else None,
                          "ustu90": sum(1 for a in l if a["gun"] > 90)})
        if sira:
            return sorted(satir, key=lambda s: sira.index(s["ad"]))
        return sorted(satir, key=lambda s: (-s["adet"], katla(s["ad"])))
    bantlar = ozet(lambda a: a["bant"], [b for b, _ in BANT])
    return {"araclar": sorted(araclar, key=lambda a: (-a["gun"], a["sasi"])), "model": ozet(lambda a: f"{a['marka']} {a['model']}".strip()), "renk": ozet(lambda a: a["renk"]),
            "bant": bantlar, "toplam": ozet(lambda a: "TOPLAM")[0], "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, bugun: date) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Yaş Bandı"
    for ad, l in (("Yaş Bandı", r["bant"] + [r["toplam"]]), ("Model", r["model"]), ("Renk", r["renk"])):
        if ad != "Yaş Bandı":
            ws = wb.create_sheet(ad)
        _baslik(ws, [ad, "Adet", "Rezerve", "Alış Değeri", "Ort. Stok Günü", "En Eski (gün)", "90+ Gün", "Finansman Maliyeti"], (24, 7, 8, 16, 10, 10, 8, 14))
        for s in l:
            n = _satir(ws, [s["ad"], s["adet"], s["rezerve"], s["alis"], s["ort_gun"], s["en_eski"], s["ustu90"], s["maliyet"]], para_s=(4, 8),
                       dolgu={7: SARI if s["ustu90"] else None})
            if s["ad"] == "TOPLAM":
                for h in ws[n]:
                    h.font = Font(bold=True)
    ws = wb.create_sheet("Araç Listesi")
    _baslik(ws, ["Şasi No", "Marka", "Model", "Versiyon", "Renk", "Model Yılı", "Stoğa Giriş", "Stok Günü", "Yaş Bandı", "Durum", "Finansman", "Faizsiz Gün", "Faizli Gün", "Oran %",
                 "Alış", "Liste", "Brüt Marj", "Finansman Maliyeti", "Marj Tüketimi %"], (19, 10, 14, 14, 12, 7, 11, 7, 11, 9, 11, 7, 7, 7, 13, 13, 12, 12, 9), "B2")
    for a in r["araclar"]:
        _satir(ws, [a["sasi"], a["marka"], a["model"], a["versiyon"], a["renk"], a["yil"], a["giris"], a["gun"], a["bant"], a["durum"], a["finansman"], a["faizsiz"], a["faizli_gun"],
                    a["oran"], a["alis"], a["liste"], a["marj"], a["maliyet"], a["marj_tuketim"]], para_s=(14, 15, 16, 17, 18, 19), tarih_s=(7,),
               dolgu={8: KIRMIZI if a["gun"] > 180 else (SARI if a["gun"] > 90 else None)})
    ws.append([])
    ws.append([f"Rapor tarihi {bugun:%d.%m.%Y}. Finansman maliyeti basit faizle, faizsiz gün düşülerek hesaplanır; gerçek faiz ve masraf için finansman sağlayıcınızın ekstresini esas alın."])
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, faiz: Decimal | None = None, firsat: Decimal | None = None, marj_esik: Decimal = Decimal(50)) -> dict:
    araclar, uy = oku(yol)
    r = analiz(araclar, bugun, faiz, firsat, marj_esik)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, bugun)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Stoktaki araçları yaşlandırıp finansman maliyetini hesaplar.")
    p.add_argument("--stok", type=Path, help="Şasi No, Marka, Model, Versiyon, Renk, Model Yılı, Stoğa Giriş Tarihi, Alış Maliyeti, Liste Fiyatı, Durum, Finansman, Faizsiz Gün")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 14.10.2026)")
    p.add_argument("--faiz", help="Floor plan / kredi yıllık faiz oranı, yüzde (örnek veride 36; verilmezse maliyet hesaplanmaz)")
    p.add_argument("--firsat-faiz", help="Özkaynakla alınan araçlar için yıllık fırsat maliyeti oranı (varsayılan: hesaplanmaz)")
    p.add_argument("--marj-esik", default="50", help="Finansman maliyeti brüt marjın bu yüzdesini geçerse uyarılır (varsayılan 50)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "arac_stok_yaslandirma.xlsx")
    a = p.parse_args(argv)
    ornek = a.stok is None
    yol = a.stok or ORNEK / "stok.csv"
    bugun = tarih(a.tarih) if a.tarih else (date(2026, 10, 14) if ornek else date.today())
    faiz = para(a.faiz) if a.faiz else (Decimal(36) if ornek else None)
    firsat = para(a.firsat_faiz) if a.firsat_faiz else None
    esik = para(a.marj_esik)
    if bugun is None or (a.faiz and (faiz is None or not 0 <= faiz <= 500)) or (a.firsat_faiz and (firsat is None or not 0 <= firsat <= 500)) or esik is None or esik <= 0:
        print("[X] --tarih GG.AA.YYYY; --faiz ve --firsat-faiz 0-500; --marj-esik 0'dan büyük olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, bugun, faiz, firsat, esik)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {t['adet']} araç ({t['rezerve']} rezerve) · alış değeri {t['alis']:,.2f} · ort. {t['ort_gun']} gün · 90+ gün {t['ustu90']} araç"
          + (f" · finansman maliyeti {t['maliyet']:,.2f}" if t["maliyet"] is not None else ""))
    for s in r["bant"]:
        print(f"     {s['ad']}: {s['adet']} araç")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
