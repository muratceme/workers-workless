import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 14)


def a(sasi, gun, alis, liste=None, model="SUV", renk="Beyaz", ozkaynak=False, faizsiz=0, durum="Stokta", yil=2026):
    return {"sasi": sasi, "marka": "M", "model": model, "versiyon": "", "renk": renk, "yil": yil, "giris": BUGUN - timedelta(days=gun), "alis": D(alis), "liste": None if liste is None else D(liste),
            "durum": durum, "finansman": "", "ozkaynak": ozkaynak, "faizsiz": faizsiz}


class AnalizTesti(unittest.TestCase):
    def test_maliyet(self):
        l = [a("A", 150, 1_000_000, 1_080_000, faizsiz=60), a("B", 200, 2_000_000, 2_100_000, "Sedan", "Gri", ozkaynak=True, yil=2025), a("C", 55, 1_000_000, 1_090_000, faizsiz=60, durum="Rezerve"),
             a("D", 10, 500_000, 480_000), a("E", 0, 800_000)]
        r = main.analiz(l, BUGUN, D(36), D(20))
        x = {v["sasi"]: v for v in r["araclar"]}
        # A: 1.000.000 × %36 × (150 − 60) ÷ 365 = 88.767,12; marj 80.000 → %110,96
        self.assertEqual((x["A"]["faizli_gun"], x["A"]["maliyet"], x["A"]["marj_tuketim"], x["A"]["bant"]), (90, D("88767.12"), D("110.96"), "91–180 gün"))
        # B özkaynak: 2.000.000 × %20 × 200 ÷ 365
        self.assertEqual((x["B"]["oran"], x["B"]["maliyet"], x["B"]["bant"]), (D(20), D("219178.08"), "180+ gün"))
        self.assertEqual((x["C"]["maliyet"], x["D"]["maliyet"], x["D"]["marj_tuketim"], x["E"]["bant"]), (D(0), D("4931.51"), None, "0–30 gün"))
        self.assertEqual([v["sasi"] for v in r["araclar"]], ["B", "A", "C", "D", "E"])
        t = r["toplam"]
        self.assertEqual((t["adet"], t["rezerve"], t["alis"], t["ort_gun"], t["en_eski"], t["ustu90"], t["maliyet"]), (5, 1, D(5_300_000), D("83.0"), 200, 2, D("312876.71")))
        self.assertEqual([(s["ad"], s["adet"]) for s in r["bant"]], [("0–30 gün", 2), ("31–60 gün", 1), ("91–180 gün", 1), ("180+ gün", 1)])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"].split(" ")[0]) for u in r["uyarilar"]),
                         [("Bilgi", "Faizsiz süre bitiyor", "C"), ("Bilgi", "Liste fiyatı maliyetin altında", "D"), ("Bilgi", "Önceki model yılı", "B"),
                          ("Orta", "90 günü geçen stok", "A"), ("Orta", "Finansman marjı eritiyor", "A"), ("Orta", "Finansman marjı eritiyor", "B"), ("Yüksek", "180 günü geçen stok", "B")])
        r = main.analiz([a("A", 100, 1000)], BUGUN)
        self.assertEqual((r["araclar"][0]["maliyet"], r["toplam"]["maliyet"], r["uyarilar"][-1]["tur"]), (None, None, "Faiz oranı verilmedi"))
        with self.assertRaises(ValueError):
            main.analiz([], BUGUN)

    def test_donusum(self):
        self.assertEqual([main._durum(x) for x in ("", "Kaporalı", "Teslim edildi", "x")], ["Stokta", "Rezerve", "Satıldı", None])
        self.assertEqual([main.bant(g) for g in (0, 30, 31, 90, 91, 181)], ["0–30 gün", "0–30 gün", "31–60 gün", "61–90 gün", "91–180 gün", "180+ gün"])
        self.assertEqual([main._ozkaynak(x) for x in ("Özkaynak", "Peşin", "Floor plan")], [True, True, False])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "stok.csv", cikti, BUGUN, D(36))
            t = r["toplam"]
            self.assertEqual((t["adet"], t["rezerve"], t["alis"], t["ort_gun"], t["ustu90"], t["maliyet"]), (59, 4, D("101751262.30"), D("59.5"), 16, D("1358763.02")))
            self.assertEqual([s["adet"] for s in r["bant"]], [26, 8, 9, 12, 4])
            self.assertEqual(sum(1 for u in r["uyarilar"] if u["tur"] == "180 günü geçen stok"), 4)
            self.assertEqual([u["tur"] for u in r["uyarilar"] if u["tur"] == "Şasi numarası biçimi"], ["Şasi numarası biçimi"])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Yaş Bandı", "Model", "Renk", "Araç Listesi", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "s.csv"
            y.write_text("VIN;Model;Giriş Tarihi;Maliyet;Durum;Faizsiz Süre\nnm0 abc;X;01.09.2026;1.000.000;;30\nNM0ABC;X;01.09.2026;1;;\nB;X;x;1;;\nC;X;01.09.2026;1;belki;\nD;Y;01.09.2026;5;Satıldı;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["sasi"], x["alis"], x["faizsiz"]) for x in l], [u["tur"] for u in uy]),
                             ([("NM0ABC", D(1000000), 30)], ["Şasi numarası biçimi", "Şasi iki kez yazılmış", "Satır okunamadı", "Satır okunamadı", "Şasi numarası biçimi"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--stok", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "14.10.2026", "--faiz", "40", "--firsat-faiz", "25", "--marj-esik", "30"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--faiz", "x"], ["--firsat-faiz", "-1"], ["--marj-esik", "0"])], [2] * 4)
            self.assertEqual(main.main(["--stok", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("VIN;Giriş Tarihi;Maliyet;Durum\nA;01.09.2026;5;Satıldı\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
