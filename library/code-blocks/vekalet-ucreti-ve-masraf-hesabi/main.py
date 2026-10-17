"""
Vekâlet Ücreti ve Masraf Hesabı — Gerisi Bende kod bloğu
Dava Takip › Hukuk Sekreteri
Avukatlık Asgari Ücret Tarifesi'nin nispi dilimleriyle (tarife dosyasından) kabul ve red edilen kısımlar için vekâlet ücretini,
yargılama giderlerinin haklılık oranına göre paylaşımını hesaplar.
  Nispi ücret      = her dilim kendi oranıyla: ilk dilim × oran1 + sonraki dilim × oran2 + … (tarifenin üçüncü kısmı)
  Maktu taban      = nispi ücret, davanın görüldüğü mahkeme için tarifedeki maktu ücretin altında kalamaz
  Üst sınır        = hükmedilen ücret kabul (ya da red) edilen miktarı geçemez
  Davacı lehine    = kabul edilen miktar üzerinden;  davalı lehine = reddedilen miktar (dava değeri − kabul) üzerinden
  Yargılama gideri = davacının yaptığı giderler × kabul oranı → davalıya yüklenir (HMK 326: haklılık oranında paylaştırma)
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
TARIFE = {"dilim": ("dilim tutari", "dilim", "tutar", "matrah"), "oran": ("oran", "oran yuzde", "yuzde")}
DOSYA = {"no": ("dosya no", "esas no", "dosya"), "deger": ("dava degeri", "deger", "talep"), "kabul": ("kabul edilen", "kabul edilen miktar", "hukmedilen"),
         "maktu": ("maktu ucret", "maktu", "maktu taban"), "harc": ("harc", "harclar", "yatirilan harc"), "bilirkisi": ("bilirkisi", "bilirkisi ucreti"),
         "tebligat": ("tebligat", "tebligat ve posta", "posta"), "tanik": ("tanik", "tanik ucreti"), "kesif": ("kesif", "kesif gideri"), "diger": ("diger", "diger masraf")}
GIDER = ("harc", "bilirkisi", "tebligat", "tanik", "kesif", "diger")
SIFIR = Decimal(0)


def tarife_oku(yol: Path) -> list[tuple[Decimal | None, Decimal]]:
    """Dilim satırları: (dilim tutarı, oran); son satırın dilim tutarı boş ya da 'kalan' olabilir (aşan kısım)."""
    l = []
    for r in kayitlar(yol, TARIFE, ("oran",)):
        if not metin(r["oran"]) and not metin(r["dilim"]):
            continue
        o = para(r["oran"])
        d = None if katla(r["dilim"]) in ("", "kalan", "asan", "asan kisim", "sonrasi") else para(r["dilim"])
        if o is None or not 0 <= o <= 100 or (katla(r["dilim"]) not in ("", "kalan", "asan", "asan kisim", "sonrasi") and (d is None or d <= 0)):
            raise ValueError(f"{yol.name} satır {r['_satir']}: dilim tutarı ya da oran okunamadı.")
        if l and l[-1][0] is None:
            raise ValueError(f"{yol.name}: 'kalan' dilimden sonra satır olamaz.")
        l.append((d, o))
    if not l:
        raise ValueError(f"{yol.name}: tarife dilimi yok.")
    return l


def nispi(tutar: Decimal, tarife: list) -> tuple[Decimal, list[dict]]:
    kalan, top, detay, alt = tutar, SIFIR, [], SIFIR
    for d, o in tarife:
        if kalan <= 0:
            break
        parca = kalan if d is None else min(kalan, d)
        u = yuvarla(parca * o / 100)
        detay.append({"alt": alt, "ust": None if d is None else alt + d, "parca": parca, "oran": o, "ucret": u})
        top += u
        kalan -= parca
        alt += parca
    if kalan > 0:
        son_oran = tarife[-1][1]
        u = yuvarla(kalan * son_oran / 100)
        detay.append({"alt": alt, "ust": None, "parca": kalan, "oran": son_oran, "ucret": u})
        top += u
    return top, detay


def vekalet(miktar: Decimal, tarife: list, maktu: Decimal) -> dict:
    if miktar <= 0:
        return {"miktar": SIFIR, "nispi": SIFIR, "ucret": SIFIR, "kural": "Ücret yok", "detay": []}
    n, detay = nispi(miktar, tarife)
    u, kural = n, "Nispi"
    if maktu and n < maktu:
        u, kural = maktu, "Maktu taban"
    if u > miktar:
        u, kural = miktar, "Miktarla sınırlı"
    return {"miktar": miktar, "nispi": n, "ucret": u, "kural": kural, "detay": detay}


def hesapla(dosyalar: list[dict], tarife: list) -> dict:
    if not dosyalar:
        raise ValueError("Hesaplanacak dosya yok.")
    sonuc, uy = [], []
    for d in dosyalar:
        kabul = d["kabul"]
        red = d["deger"] - kabul
        oran = kabul / d["deger"] if d["deger"] else SIFIR
        gider = sum((d[g] for g in GIDER), SIFIR)
        s = {**d, "red": red, "oran": yuvarla(oran * 100), "davaci": vekalet(kabul, tarife, d["maktu"]), "davali": vekalet(red, tarife, d["maktu"]), "gider": gider,
             "gider_davali": yuvarla(gider * oran), "gider_davaci": gider - yuvarla(gider * oran)}
        if not d["maktu"]:
            uy.append(uyari("Orta", "Maktu ücret yazılmadı", d["no"], "Mahkemeye göre maktu ücret girilmediği için maktu taban kontrolü yapılmadı."))
        for taraf in ("davaci", "davali"):
            if s[taraf]["kural"] == "Miktarla sınırlı":
                uy.append(uyari("Bilgi", "Ücret miktarla sınırlandı", d["no"], f"{'Davacı' if taraf == 'davaci' else 'Davalı'} lehine ücret, {'kabul' if taraf == 'davaci' else 'red'} edilen {s[taraf]['miktar']:,.2f} TL'yi geçemediği için bu tutara indirildi."))
        sonuc.append(s)
    return {"dosyalar": sonuc, "uyarilar": uy}


def dosya_oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, DOSYA, ("deger", "kabul")):
        no = metin(r["no"]) or f"Satır {r['_satir']}"
        if not metin(r["deger"]) and not metin(r["kabul"]):
            continue
        v = {a: (para(r[a]) if metin(r[a]) else SIFIR) for a in ("deger", "kabul", "maktu") + GIDER}
        if any(x is None or x < 0 for x in v.values()) or v["deger"] <= 0 or v["kabul"] > v["deger"]:
            uy.append(uyari("Yüksek", "Satır okunamadı", no, f"Satır {r['_satir']}: tutarlardan biri okunamadı, negatif ya da kabul edilen miktar dava değerinden büyük."))
            continue
        l.append({"no": no, **v})
    return l, uy


def rapor_yaz(cikti: Path, r: dict, tarife: list) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Dosyalar"
    _baslik(ws, ["Dosya", "Dava Değeri", "Kabul", "Red", "Kabul Oranı %", "Maktu Taban", "Davacı Lehine Nispi", "Davacı Lehine Ücret", "Kural", "Davalı Lehine Nispi",
                 "Davalı Lehine Ücret", "Kural", "Toplam Gider", "Davalıya Yüklenen Gider", "Davacıda Kalan Gider"], (14, 14, 14, 14, 9, 12, 13, 13, 15, 13, 13, 15, 12, 13, 13), "B2")
    for s in r["dosyalar"]:
        _satir(ws, [s["no"], s["deger"], s["kabul"], s["red"], s["oran"], s["maktu"] or None, s["davaci"]["nispi"], s["davaci"]["ucret"], s["davaci"]["kural"], s["davali"]["nispi"],
                    s["davali"]["ucret"], s["davali"]["kural"], s["gider"], s["gider_davali"], s["gider_davaci"]], para_s=(2, 3, 4, 5, 6, 7, 8, 10, 11, 13, 14, 15))
    ws.append([])
    ws.append(["Tutarlara KDV dahil değildir. Nispi ücret tarife dosyasındaki dilimlerle hesaplanır; maktu taban ve miktar sınırı uygulanır. Sonuç avukat tarafından kontrol edilmelidir."])
    ws = wb.create_sheet("Dilim Detayı")
    _baslik(ws, ["Dosya", "Taraf", "Dilim Alt", "Dilim Üst", "Dilime Düşen", "Oran %", "Ücret"], (14, 10, 14, 14, 14, 8, 13))
    for s in r["dosyalar"]:
        for taraf, ad in (("davaci", "Davacı"), ("davali", "Davalı")):
            for d in s[taraf]["detay"]:
                _satir(ws, [s["no"], ad, d["alt"], d["ust"], d["parca"], d["oran"], d["ucret"]], para_s=(3, 4, 5, 6, 7))
    ws = wb.create_sheet("Tarife")
    _baslik(ws, ["Dilim Tutarı", "Oran %"], (16, 8))
    for d, o in tarife:
        _satir(ws, [d if d is not None else "Kalan", o], para_s=(1, 2))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(tarife_yol: Path, dosyalar: list[dict], cikti: Path, uy: list[dict] | None = None) -> dict:
    tarife = tarife_oku(tarife_yol)
    r = hesapla(dosyalar, tarife)
    r["uyarilar"] = sirala((uy or []) + r["uyarilar"])
    rapor_yaz(cikti, r, tarife)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="AAÜT nispi dilimleriyle vekâlet ücretini ve yargılama giderlerinin paylaşımını hesaplar.")
    p.add_argument("--tarife", type=Path, help="Dilim Tutarı, Oran %% (son satır 'Kalan'); varsayılan: örnek 2025-2026 AAÜT dilimleri")
    p.add_argument("--dosyalar", type=Path, help="Dosya No, Dava Değeri, Kabul Edilen, Maktu Ücret, Harç, Bilirkişi, Tebligat, Tanık, Keşif, Diğer")
    p.add_argument("--deger", help="Tek dosya: dava değeri (TL)")
    p.add_argument("--kabul", help="Tek dosya: kabul edilen miktar (varsayılan dava değerinin tamamı)")
    p.add_argument("--maktu", default="0", help="Tek dosya: mahkemeye göre tarifedeki maktu ücret (TL)")
    p.add_argument("--gider", default="0", help="Tek dosya: davacının yaptığı toplam yargılama gideri (TL)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "vekalet_ucreti.xlsx")
    a = p.parse_args(argv)
    tarife = a.tarife or ORNEK / "tarife_2025_2026.csv"
    uy = []
    if a.deger:
        dg, kb, mk, gd = para(a.deger), para(a.kabul) if a.kabul else para(a.deger), para(a.maktu), para(a.gider)
        if dg is None or dg <= 0 or kb is None or not 0 <= kb <= dg or mk is None or mk < 0 or gd is None or gd < 0:
            print("[X] --deger 0'dan büyük; --kabul 0 ile dava değeri arası; --maktu ve --gider 0 ya da üstü olmalı")
            return 2
        dosyalar = [{"no": "Dosya", "deger": dg, "kabul": kb, "maktu": mk, "harc": SIFIR, "bilirkisi": SIFIR, "tebligat": SIFIR, "tanik": SIFIR, "kesif": SIFIR, "diger": gd}]
    else:
        yol = a.dosyalar or ORNEK / "dosyalar.csv"
        if not yol.exists():
            print(f"[X] Dosya bulunamadı: {yol}")
            return 1
        try:
            dosyalar, uy = dosya_oku(yol)
        except ValueError as h:
            print(f"[X] {h}")
            return 1
    if not tarife.exists():
        print(f"[X] Dosya bulunamadı: {tarife}")
        return 1
    try:
        r = calistir(tarife, dosyalar, a.cikti, uy)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    for s in r["dosyalar"]:
        print(f"[OK] {s['no']}: değer {s['deger']:,.2f} · kabul {s['kabul']:,.2f} (%{s['oran']}) · davacı lehine {s['davaci']['ucret']:,.2f} ({s['davaci']['kural']}) · "
              f"davalı lehine {s['davali']['ucret']:,.2f} ({s['davali']['kural']}) · davalıya gider {s['gider_davali']:,.2f}")
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
