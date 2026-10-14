import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
TARIFE = main.tarife_oku(ORNEK / "tarife_2025_2026.csv")


def dosya(deger, kabul, maktu=45000, **gider):
    return {"no": "X", "deger": D(deger), "kabul": D(kabul), "maktu": D(maktu), **{g: D(gider.get(g, 0)) for g in main.GIDER}}


class HesapTesti(unittest.TestCase):
    def test_nispi(self):
        # Tarife özetlerindeki örnekler: 500.000 → 80.000; 900.000 → 96.000 + 45.000 = 141.000
        self.assertEqual([main.nispi(D(x), TARIFE)[0] for x in (500000, 900000, 600000, 18600000, 25000000)], [D(80000), D(141000), D(96000), D(1242000), D(1306000)])
        self.assertEqual(len(TARIFE), 10)
        self.assertEqual([d["oran"] for d in main.nispi(D(1300000), TARIFE)[1]], [D(16), D(15), D(14)])

    def test_vekalet(self):
        self.assertEqual([main.vekalet(D(x), TARIFE, D(45000))["ucret"] for x in (100000, 900000, 30000, 0)], [D(45000), D(141000), D(30000), D(0)])
        self.assertEqual([main.vekalet(D(x), TARIFE, D(45000))["kural"] for x in (100000, 900000, 30000, 0)], ["Maktu taban", "Nispi", "Miktarla sınırlı", "Ücret yok"])
        self.assertEqual(main.vekalet(D(100000), TARIFE, D(0))["ucret"], D(16000))

    def test_paylasim(self):
        r = main.hesapla([dosya(2500000, 1600000, harc=40000, bilirkisi=20000), dosya(100000, 100000, 0)], TARIFE)
        s = r["dosyalar"][0]
        self.assertEqual((s["red"], s["oran"], s["davaci"]["ucret"], s["davali"]["ucret"], s["gider"], s["gider_davali"], s["gider_davaci"]),
                         (D(900000), D("64.00"), D(242000), D(141000), D(60000), D(38400), D(21600)))
        self.assertEqual([u["tur"] for u in r["uyarilar"]], ["Maktu ücret yazılmadı"])
        with self.assertRaises(ValueError):
            main.hesapla([], TARIFE)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            l, uy = main.dosya_oku(ORNEK / "dosyalar.csv")
            r = main.calistir(ORNEK / "tarife_2025_2026.csv", l, cikti, uy)
            self.assertEqual([(s["davaci"]["ucret"], s["davali"]["ucret"]) for s in r["dosyalar"]],
                             [(D(141000), D(0)), (D(45000), D(0)), (D(242000), D(141000)), (D(30000), D(10000)), (D(1306000), D(0)), (D(0), D(118500)), (D(126000), D(64000))])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Dosyalar", "Dilim Detayı", "Tarife", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            t, d = Path(tmp) / "t.csv", Path(tmp) / "d.csv"
            t.write_text("Dilim;Oran\n100.000;10\nKalan;5\n", encoding="utf-8")
            self.assertEqual(main.tarife_oku(t), [(D(100000), D(10)), (None, D(5))])
            self.assertEqual(main.nispi(D(300000), main.tarife_oku(t))[0], D(20000))
            t2 = Path(tmp) / "t2.csv"
            t2.write_text("Dilim;Oran\n100.000;10\n", encoding="utf-8")
            self.assertEqual(main.nispi(D(300000), main.tarife_oku(t2))[0], D(30000))
            for bozuk in ("Dilim;Oran\nKalan;5\n100;3\n", "Dilim;Oran\nx;5\n", "Dilim;Oran\n100;150\n", "Dilim;Oran\n"):
                t2.write_text(bozuk, encoding="utf-8")
                with self.assertRaises(ValueError):
                    main.tarife_oku(t2)
            d.write_text("Esas No;Dava Değeri;Kabul Edilen;Maktu\n1;100.000;50.000;45.000\n2;100;200;\n3;x;1;\n", encoding="utf-8")
            l, uy = main.dosya_oku(d)
            self.assertEqual((len(l), len(uy)), (1, 2))
            c = str(Path(tmp) / "x.xlsx")
            self.assertEqual(main.main(["--deger", "900.000", "--maktu", "45.000", "--gider", "1.000", "--cikti", c]), 0)
            self.assertEqual(main.main(["--dosyalar", str(d), "--tarife", str(t), "--cikti", c]), 0)
            self.assertEqual([main.main(v + ["--cikti", c]) for v in (["--deger", "0"], ["--deger", "10", "--kabul", "20"], ["--deger", "10", "--maktu", "x"])], [2] * 3)
            self.assertEqual(main.main(["--dosyalar", str(Path(tmp) / "yok.csv")]), 1)
            self.assertEqual(main.main(["--deger", "10", "--tarife", str(Path(tmp) / "yok.csv")]), 1)
            self.assertEqual(main.main(["--deger", "10", "--tarife", str(t2), "--cikti", c]), 1)
            d.write_text("Esas No;Dava Değeri;Kabul Edilen\n1;x;1\n", encoding="utf-8")
            self.assertEqual(main.main(["--dosyalar", str(d), "--cikti", c]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
