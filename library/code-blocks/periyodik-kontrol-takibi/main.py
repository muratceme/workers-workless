"""
Periyodik Kontrol Takibi — Gerisi Bende kod bloğu
İş Sağlığı ve Güvenliği › İş Güvenliği Uzmanı

Kaldırma ekipmanı, basınçlı kap, elektrik tesisatı ve yangın ekipmanı gibi iş ekipmanlarının periyodik kontrol tarihlerini izler.
  Sonraki kontrol   = son kontrol tarihi + kontrol periyodu (ay)   (ay sonu taşarsa ayın son günü)
  Kalan gün         = sonraki kontrol − durum tarihi
  Durum             = Süresi geçmiş (kalan < 0) · Yaklaşıyor (kalan ≤ 30 gün) · Uygun
                      Hiç kontrol edilmemiş (son kontrol tarihi yok) · Periyot tanımsız (periyot yok)
  Son kontrol sonucu "Uygun Değil" ise tarih ne olursa olsun ayrıca uyarılır.
Kontrol periyodunu kod bilmez; her ekipman için siz yazarsınız. Süreler mevzuata, ilgili standarda ve imalatçı talimatına göre değişir.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
import calendar

SUTUN = {"no": ("ekipman no", "ekipman kodu", "seri no", "no"), "ad": ("ekipman", "ekipman adi", "tanim"), "tur": ("tur", "ekipman turu", "grup"), "konum": ("konum", "yer", "bolum"),
         "son": ("son kontrol tarihi", "son kontrol", "kontrol tarihi"), "periyot": ("kontrol periyodu ay", "periyot ay", "kontrol periyodu", "periyot"), "sonuc": ("sonuc", "kontrol sonucu", "son kontrol sonucu"),
         "yapan": ("kontrolu yapan", "yetkili kisi", "kontrol firmasi", "firma"), "rapor": ("rapor no", "rapor numarasi")}
UYGUN_DEGIL = {"uygun degil", "uygunsuz", "kusurlu", "kullanilamaz", "sakincali", "uygun degildir"}
DURUMLAR = ("Süresi geçmiş", "Hiç kontrol edilmemiş", "Yaklaşıyor", "Periyot tanımsız", "Uygun")
DURUM_RENK = {"Süresi geçmiş": KIRMIZI, "Hiç kontrol edilmemiş": KIRMIZI, "Yaklaşıyor": SARI, "Periyot tanımsız": MAVI, "Uygun": YESIL}


def ay_ekle(t: date, n: int) -> date:
    y, a = divmod(t.year * 12 + t.month - 1 + n, 12)
    return date(y, a + 1, min(t.day, calendar.monthrange(y, a + 1)[1]))


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("ad",)):
        ad, no = metin(r["ad"]), metin(r["no"])
        if not ad:
            continue
        kim = f"{no} · {ad}".strip(" ·")
        son, periyot = tarih(r["son"]), None
        if metin(r["son"]) and son is None:
            uy.append(uyari("Orta", "Tarih okunamadı", kim, f"Son kontrol tarihi '{metin(r['son'])}' okunamadı; boş sayıldı."))
        if metin(r["periyot"]):
            m = re.fullmatch(r"\s*(\d{1,3})\s*(ay)?\s*", metin(r["periyot"]).lower())
            periyot = int(m.group(1)) if m and int(m.group(1)) > 0 else None
            if periyot is None:
                uy.append(uyari("Orta", "Periyot okunamadı", kim, f"Kontrol periyodu '{metin(r['periyot'])}' okunamadı; ay olarak tam sayı yazın (örn. 12)."))
        s = katla(r["sonuc"])
        l.append({"no": no, "ad": ad, "tur": metin(r["tur"]) or "(tür yok)", "konum": metin(r["konum"]) or "(konum yok)", "son": son, "periyot": periyot, "sonuc": "Uygun Değil" if s in UYGUN_DEGIL else metin(r["sonuc"]),
                  "yapan": metin(r["yapan"]), "rapor": metin(r["rapor"])})
    return l, uy


def degerlendir(ekipman: list[dict], bugun: date, gun: int = 30) -> dict:
    if not ekipman:
        raise ValueError("Değerlendirilecek ekipman yok.")
    uy = []
    for e in ekipman:
        kim = f"{e['no']} · {e['ad']}".strip(" ·")
        e["sonraki"] = ay_ekle(e["son"], e["periyot"]) if e["son"] and e["periyot"] else None
        e["kalan"] = (e["sonraki"] - bugun).days if e["sonraki"] else None
        if e["son"] is None:
            e["durum"] = "Hiç kontrol edilmemiş"
            uy.append(uyari("Yüksek", "Hiç kontrol edilmemiş", kim, f"Son kontrol tarihi yok. Konum: {e['konum']}. Kontrol yapılmadıysa ekipmanın kullanımını gözden geçirin; yapıldıysa tarihi ve rapor numarasını girin."))
        elif e["periyot"] is None:
            e["durum"] = "Periyot tanımsız"
            uy.append(uyari("Orta", "Periyot tanımsız", kim, f"Son kontrol {e['son']:%d.%m.%Y} ama kontrol periyodu yazılmamış; sonraki kontrol tarihi hesaplanamadı."))
        elif e["kalan"] < 0:
            e["durum"] = "Süresi geçmiş"
            uy.append(uyari("Yüksek", "Kontrol süresi geçmiş", kim, f"Son kontrol {e['son']:%d.%m.%Y}, periyot {e['periyot']} ay; kontrol {e['sonraki']:%d.%m.%Y} tarihinde yapılmalıydı ({-e['kalan']} gün geçti). Konum: {e['konum']}."))
        elif e["kalan"] <= gun:
            e["durum"] = "Yaklaşıyor"
            uy.append(uyari("Orta", "Kontrol yaklaşıyor", kim, f"Sonraki kontrol {e['sonraki']:%d.%m.%Y} ({e['kalan']} gün kaldı). Konum: {e['konum']}."))
        else:
            e["durum"] = "Uygun"
        if e["sonuc"] == "Uygun Değil":
            uy.append(uyari("Yüksek", "Son kontrolde uygun bulunmadı", kim, f"Son kontrolün sonucu 'Uygun Değil'" + (f" (rapor {e['rapor']})" if e["rapor"] else "") + ". Eksiklik giderilip yeniden kontrol edilene kadar kullanımını değerlendirin."))
        if e["son"] and e["son"] > bugun:
            uy.append(uyari("Orta", "Kontrol tarihi ileride", kim, f"Son kontrol tarihi {e['son']:%d.%m.%Y}, durum tarihi {bugun:%d.%m.%Y}."))

    def ozet(alan: str) -> list[dict]:
        grup = {}
        for e in ekipman:
            g = grup.setdefault(katla(e[alan]), {"ad": e[alan], "adet": 0, **{d: 0 for d in DURUMLAR}, "uygun_degil": 0})
            g["adet"] += 1
            g[e["durum"]] += 1
            g["uygun_degil"] += e["sonuc"] == "Uygun Değil"
        return [g for _, g in sorted(grup.items())]

    isler = sorted((e for e in ekipman if e["durum"] in ("Süresi geçmiş", "Hiç kontrol edilmemiş", "Yaklaşıyor")), key=lambda e: (e["sonraki"] or date.min, katla(e["no"]), katla(e["ad"])))
    plan = defaultdict(int)
    for e in ekipman:
        if e["durum"] in ("Süresi geçmiş", "Hiç kontrol edilmemiş"):
            plan["Gecikmiş / hemen"] += 1
        elif e["sonraki"]:
            plan[f"{e['sonraki']:%Y-%m}"] += 1
    return {"ekipman": ekipman, "say": {d: sum(1 for e in ekipman if e["durum"] == d) for d in DURUMLAR}, "uygun_degil": sum(1 for e in ekipman if e["sonuc"] == "Uygun Değil"), "isler": isler, "turler": ozet("tur"), "konumlar": ozet("konum"),
            "plan": sorted(plan.items(), key=lambda p: (p[0] != "Gecikmiş / hemen", p[0])), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, bugun: date) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ekipmanlar"
    basliklar = ["Durum", "Ekipman No", "Ekipman", "Tür", "Konum", "Son Kontrol", "Periyot (ay)", "Sonraki Kontrol", "Kalan Gün", "Son Kontrol Sonucu", "Kontrolü Yapan", "Rapor No"]
    _baslik(ws, basliklar, (22, 13, 32, 22, 20, 12, 9, 13, 9, 14, 24, 14))
    for e in r["ekipman"]:
        _satir(ws, [e["durum"], e["no"], e["ad"], e["tur"], e["konum"], e["son"], e["periyot"], e["sonraki"], e["kalan"], e["sonuc"], e["yapan"], e["rapor"]], tarih_s=(6, 8),
               dolgu={1: DURUM_RENK[e["durum"]], 10: KIRMIZI if e["sonuc"] == "Uygun Değil" else None})
    ws = wb.create_sheet("Yapılacak Kontroller")
    _baslik(ws, basliklar[:9] + ["Planlanan Tarih (siz doldurun)"], (22, 13, 32, 22, 20, 12, 9, 13, 9, 20))
    for e in r["isler"]:
        _satir(ws, [e["durum"], e["no"], e["ad"], e["tur"], e["konum"], e["son"], e["periyot"], e["sonraki"], e["kalan"], None], tarih_s=(6, 8), dolgu={1: DURUM_RENK[e["durum"]]})
    for baslik, alan, ad in (("Tür Özeti", "turler", "Tür"), ("Konum Özeti", "konumlar", "Konum")):
        ws = wb.create_sheet(baslik)
        _baslik(ws, [ad, "Ekipman"] + list(DURUMLAR) + ["Son Sonucu Uygun Değil"], (26, 9, 11, 13, 11, 11, 9, 12))
        for g in r[alan]:
            _satir(ws, [g["ad"], g["adet"]] + [g[d] or None for d in DURUMLAR] + [g["uygun_degil"] or None], dolgu={3: KIRMIZI if g["Süresi geçmiş"] else None, 4: KIRMIZI if g["Hiç kontrol edilmemiş"] else None})
    ws = wb.create_sheet("Aylık Plan")
    _baslik(ws, ["Ay", "Kontrol Sayısı"], (20, 13))
    ws.append([f"Durum tarihi: {bugun:%d.%m.%Y}"])
    for ay, n in r["plan"]:
        _satir(ws, [ay, n], dolgu={1: KIRMIZI if ay.startswith("Gecikmiş") else None})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, gun: int = 30) -> dict:
    ekipman, uy = oku(yol)
    r = degerlendir(ekipman, bugun, gun)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, bugun)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="İş ekipmanlarının periyodik kontrol tarihlerini izler.")
    p.add_argument("--ekipman", type=Path, help="Ekipman No, Ekipman, Tür, Konum, Son Kontrol Tarihi, Kontrol Periyodu (ay), Sonuç, Kontrolü Yapan, Rapor No")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--gun", type=int, default=30, help="Kaç gün önceden uyarılsın (varsayılan 30)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "periyodik_kontrol_takibi.xlsx")
    a = p.parse_args(argv)
    ornek = a.ekipman is None
    yol = ORNEK / "ekipman.csv" if ornek else a.ekipman
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    if bugun is None or not 0 <= a.gun <= 365:
        print("[X] --tarih GG.AA.YYYY; --gun 0–365 arasında olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, bugun, a.gun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    s = r["say"]
    print(f"[OK] {len(r['ekipman'])} ekipman · süresi geçmiş {s['Süresi geçmiş']}, hiç kontrol edilmemiş {s['Hiç kontrol edilmemiş']}, yaklaşıyor {s['Yaklaşıyor']}, uygun {s['Uygun']}" + (f" · son sonucu uygun değil {r['uygun_degil']}" if r["uygun_degil"] else ""))
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
