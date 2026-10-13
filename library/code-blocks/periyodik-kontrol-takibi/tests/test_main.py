import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)


def e(no, son, periyot=12, sonuc="Uygun", tur="T", konum="K"):
    return {"no": no, "ad": "Ad", "tur": tur, "konum": konum, "son": son, "periyot": periyot, "sonuc": sonuc, "yapan": "", "rapor": "R1"}


class DegerlendirTesti(unittest.TestCase):
    def test_ay_ekle(self):
        self.assertEqual([main.ay_ekle(date(*t), n) for t, n in (((2026, 7, 1), 3), ((2026, 8, 31), 6), ((2027, 8, 31), 6), ((2026, 11, 30), 3), ((2026, 12, 15), 1), ((2026, 1, 31), 1), ((2026, 5, 31), 120))],
                         [date(2026, 10, 1), date(2027, 2, 28), date(2028, 2, 29), date(2027, 2, 28), date(2027, 1, 15), date(2026, 2, 28), date(2036, 5, 31)])

    def test_durumlar(self):
        l = [e("dün", date(2025, 10, 11)), e("bugün", date(2025, 10, 12)), e("otuz", date(2025, 11, 11)), e("otuzbir", date(2025, 11, 12)), e("hiç", None), e("periyotsuz", date(2026, 1, 1), None),
             e("kusurlu", date(2026, 9, 1), sonuc="Uygun Değil"), e("ileri", date(2026, 11, 1), 3), e("üç", date(2026, 7, 13), 3), e("hiç2", None, None)]
        r = main.degerlendir(l, BUGUN)
        self.assertEqual([(x["no"], x["durum"], x["sonraki"], x["kalan"]) for x in r["ekipman"]],
                         [("dün", "Süresi geçmiş", date(2026, 10, 11), -1), ("bugün", "Yaklaşıyor", date(2026, 10, 12), 0), ("otuz", "Yaklaşıyor", date(2026, 11, 11), 30), ("otuzbir", "Uygun", date(2026, 11, 12), 31),
                          ("hiç", "Hiç kontrol edilmemiş", None, None), ("periyotsuz", "Periyot tanımsız", None, None), ("kusurlu", "Uygun", date(2027, 9, 1), 324), ("ileri", "Uygun", date(2027, 2, 1), 112),
                          ("üç", "Yaklaşıyor", date(2026, 10, 13), 1), ("hiç2", "Hiç kontrol edilmemiş", None, None)])
        self.assertEqual((r["say"], r["uygun_degil"]), ({"Süresi geçmiş": 1, "Hiç kontrol edilmemiş": 2, "Yaklaşıyor": 3, "Periyot tanımsız": 1, "Uygun": 3}, 1))
        self.assertEqual([x["no"] for x in r["isler"]], ["hiç", "hiç2", "dün", "bugün", "üç", "otuz"])
        self.assertEqual(r["plan"], [("Gecikmiş / hemen", 3), ("2026-10", 2), ("2026-11", 2), ("2027-02", 1), ("2027-09", 1)])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         sorted([("Yüksek", "Kontrol süresi geçmiş", "dün · Ad"), ("Orta", "Kontrol yaklaşıyor", "bugün · Ad"), ("Orta", "Kontrol yaklaşıyor", "otuz · Ad"), ("Yüksek", "Hiç kontrol edilmemiş", "hiç · Ad"),
                                 ("Orta", "Periyot tanımsız", "periyotsuz · Ad"), ("Yüksek", "Son kontrolde uygun bulunmadı", "kusurlu · Ad"), ("Orta", "Kontrol tarihi ileride", "ileri · Ad"), ("Orta", "Kontrol yaklaşıyor", "üç · Ad"),
                                 ("Yüksek", "Hiç kontrol edilmemiş", "hiç2 · Ad")]))
        d = lambda gun: [x["durum"] for x in main.degerlendir([e("a", date(2025, 11, 26)), e("b", date(2025, 11, 27))], BUGUN, gun)["ekipman"]]   # noqa: E731
        self.assertEqual((d(45), d(46), d(0)), (["Yaklaşıyor", "Uygun"], ["Yaklaşıyor", "Yaklaşıyor"], ["Uygun", "Uygun"]))
        with self.assertRaises(ValueError):
            main.degerlendir([], BUGUN)

    def test_ozet(self):
        l = [e("1", date(2025, 1, 1), tur="Vinç", konum="A"), e("2", date(2026, 9, 1), tur="vinç", konum="B", sonuc="Uygun Değil"), e("3", None, tur="Kazan", konum="A"), e("4", date(2025, 11, 1), tur="Kazan", konum="a")]
        r = main.degerlendir(l, BUGUN)
        k = lambda liste: [(g["ad"], g["adet"], g["Süresi geçmiş"], g["Hiç kontrol edilmemiş"], g["Yaklaşıyor"], g["Uygun"], g["uygun_degil"]) for g in liste]   # noqa: E731
        self.assertEqual((k(r["turler"]), k(r["konumlar"])), ([("Kazan", 2, 0, 1, 1, 0, 0), ("Vinç", 2, 1, 0, 0, 1, 1)], [("A", 3, 1, 1, 1, 0, 0), ("B", 1, 0, 0, 0, 1, 1)]))


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "ekipman.csv", cikti, BUGUN)
            self.assertEqual([(x["no"], x["durum"], x["sonraki"], x["kalan"]) for x in r["ekipman"]],
                             [("PK-001", "Yaklaşıyor", date(2026, 10, 20), 8), ("PK-002", "Süresi geçmiş", date(2026, 10, 1), -11), ("PK-003", "Uygun", date(2026, 12, 15), 64), ("PK-004", "Yaklaşıyor", date(2026, 11, 5), 24),
                              ("PK-005", "Uygun", date(2027, 3, 18), 157), ("PK-006", "Süresi geçmiş", date(2026, 10, 10), -2), ("PK-007", "Süresi geçmiş", date(2026, 10, 10), -2), ("PK-008", "Uygun", date(2027, 4, 22), 192),
                              ("PK-009", "Uygun", date(2027, 1, 30), 110), ("PK-010", "Hiç kontrol edilmemiş", None, None), ("PK-011", "Periyot tanımsız", None, None), ("PK-012", "Uygun", date(2027, 8, 31), 323)])
            self.assertEqual((r["say"], r["uygun_degil"]), ({"Süresi geçmiş": 3, "Hiç kontrol edilmemiş": 1, "Yaklaşıyor": 2, "Periyot tanımsız": 1, "Uygun": 5}, 1))
            self.assertEqual([x["no"] for x in r["isler"]], ["PK-010", "PK-002", "PK-006", "PK-007", "PK-001", "PK-004"])
            self.assertEqual([(g["ad"], g["adet"], g["Süresi geçmiş"], g["Hiç kontrol edilmemiş"], g["Yaklaşıyor"], g["Periyot tanımsız"], g["Uygun"]) for g in r["turler"]],
                             [("Basınçlı Kap", 2, 0, 0, 1, 0, 1), ("Depolama", 1, 0, 0, 0, 0, 1), ("Kaldırma ve İletme", 3, 1, 0, 1, 0, 1), ("Tesisat", 4, 2, 0, 0, 1, 1), ("Yangın", 2, 0, 1, 0, 0, 1)])
            self.assertEqual(r["plan"], [("Gecikmiş / hemen", 4), ("2026-10", 1), ("2026-11", 1), ("2026-12", 1), ("2027-01", 1), ("2027-03", 1), ("2027-04", 1), ("2027-08", 1)])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Yüksek", "Hiç kontrol edilmemiş", "PK-010 · Yangın Söndürme Cihazları"), ("Yüksek", "Kontrol süresi geçmiş", "PK-002 · Tavan Vinci 5 ton"), ("Yüksek", "Kontrol süresi geçmiş", "PK-006 · Elektrik Tesisatı"),
                              ("Yüksek", "Kontrol süresi geçmiş", "PK-007 · Topraklama Tesisatı"), ("Yüksek", "Son kontrolde uygun bulunmadı", "PK-007 · Topraklama Tesisatı"), ("Orta", "Kontrol yaklaşıyor", "PK-001 · Forklift 2,5 ton"),
                              ("Orta", "Kontrol yaklaşıyor", "PK-004 · Hava Kompresörü Tankı"), ("Orta", "Periyot tanımsız", "PK-011 · Havalandırma Tesisatı")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Ekipmanlar", "Yapılacak Kontroller", "Tür Özeti", "Konum Özeti", "Aylık Plan", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "e.csv"
            y.write_text("Seri No;Ekipman Adı;Grup;Bölüm;Son Kontrol;Periyot;Kontrol Sonucu;Firma;Rapor No\nS1;Vinç;;;01.07.2026;3 Ay;KUSURLU;F;R\nS2;Kazan;B;K;x;0;Uygun;;\n;Tank;;;2026-01-05;12;;;\nS4;Lift;;;;yılda bir;;;\n"
                         "S5;;;;;;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["ad"], x["tur"], x["konum"], x["son"], x["periyot"], x["sonuc"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("S1", "Vinç", "(tür yok)", "(konum yok)", date(2026, 7, 1), 3, "Uygun Değil"), ("S2", "Kazan", "B", "K", None, None, "Uygun"), ("", "Tank", "(tür yok)", "(konum yok)", date(2026, 1, 5), 12, ""),
                               ("S4", "Lift", "(tür yok)", "(konum yok)", None, None, "")], [("Tarih okunamadı", "S2 · Kazan"), ("Periyot okunamadı", "S2 · Kazan"), ("Periyot okunamadı", "S4 · Lift")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--ekipman", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "12.10.2026", "--gun", "45"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + x) for x in (["--tarih", "x"], ["--gun", "-1"], ["--gun", "366"])], [2] * 3)
            self.assertEqual(main.main(["--ekipman", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Tür;Konum\nA;1\n", "Ekipman;Konum\n;1\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
