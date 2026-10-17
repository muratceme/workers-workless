"""
Faiz Riski Vade Uyumsuzluk (Gap) Analizi — Gerisi Bende kod bloğu
Hazine ve Aktif-Pasif Yönetimi › Aktif-Pasif Yönetimi Uzmanı

Faize duyarlı varlık ve yükümlülükleri yeniden fiyatlama vadelerine göre dilimlere dağıtır; gap tablosunu ve paralel
faiz şokunun net faiz gelirine ve ekonomik değere yaklaşık etkisini hesaplar.
  Yeniden fiyatlama günü : sabit faizli kalemde vade tarihi; değişken faizlide bir sonraki faiz belirleme tarihi.
  Süre t (yıl)           : (yeniden fiyatlama günü − rapor tarihi) ÷ 365; geçmiş tarih 0 sayılır.
  Gap (dilim)            : faize duyarlı varlıklar − faize duyarlı yükümlülükler;  kümülatif gap = dilimlerin birikimli toplamı.
  Net faiz geliri etkisi (12 ay), kalem bazında:  ΔNFG = ± tutar × Δr × (1 − t)   yalnız t < 1 yıl için
                          (kalem t anında yeni orandan fiyatlanır, yılın kalanında Δr kadar fazla / eksik faiz işler;
                          varlık +, yükümlülük −). Bilanço sabit, şok anlık ve paralel varsayılır.
  Ekonomik değer etkisi  : ΔED ≈ ∓ tutar × t ÷ (1 + faiz oranı) × Δr   (varlık −, yükümlülük +; t düzeltilmiş süre yerine
                          kullanılan yaklaşık değerdir; ara kupon ve taksitler yok sayılır → uzun vadede etkiyi büyütür).
  Faize duyarsız kalemler (nakit, özkaynak, faizsiz mevduat...) gap'e girmez, ayrı gösterilir.
Şok büyüklükleri --sok ile baz puan olarak verilir (varsayılan +100, −100). Yasal ölçüm için BDDK'nın bankacılık
hesaplarından kaynaklanan faiz oranı riskine ilişkin düzenlemesindeki dilim ve şoklar esastır. İnternete bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: 22 kalem, TL
    python main.py --kalemler bilanco.xlsx --tarih 30.09.2026 --sok 100,-100,300 --ozkaynak 2.400.000.000
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUNLAR = {"kalem": ("kalem", "hesap", "aciklama"), "tur": ("tur", "varlik yukumluluk", "taraf"), "tutar": ("tutar", "bakiye", "tutar tl"),
            "tip": ("faiz tipi", "faiz turu", "tip"), "vade": ("vade tarihi", "vade"), "yf": ("yeniden fiyatlama tarihi", "yeniden fiyatlama", "sonraki faiz tarihi"),
            "oran": ("faiz orani", "oran", "faiz orani yuzde"), "pb": ("para birimi", "pb")}
# (ad, üst sınır gün)
DILIMLER = [("0-1 ay", 30), ("1-3 ay", 91), ("3-6 ay", 182), ("6-12 ay", 365), ("1-2 yıl", 730), ("2-5 yıl", 1826), ("5 yıl+", None)]
DUYARSIZ = "Faize duyarsız"


def bin_(x) -> str:
    return f"{x:,.0f}".replace(",", ".")


def dilim_bul(gun: int) -> str:
    return next((ad for ad, ust in DILIMLER if ust is None or gun <= ust), DILIMLER[-1][0])


def oku(yol: Path, rapor: date, pb_filtre: str = "") -> tuple[list[dict], list[dict]]:
    uy, kalemler = [], []
    for r in kayitlar(yol, SUTUNLAR, ("kalem", "tur", "tutar")):
        ad = metin(r["kalem"])
        if not ad:
            continue
        pb = metin(r["pb"]).upper().replace("TRY", "TL") or "TL"
        if pb_filtre and pb != pb_filtre:
            continue
        tur_k, t = katla(r["tur"]), para(r["tutar"])
        tur = "Varlık" if tur_k.startswith(("varlik", "aktif")) else "Yükümlülük" if tur_k.startswith(("yukumluluk", "pasif", "kaynak", "ozkaynak")) else None
        if tur is None or t is None or t < 0:
            uy.append(uyari("Yüksek", "Satır okunamadı", ad, f"Satır {r['_satir']}: tür '{metin(r['tur'])}' (Varlık / Yükümlülük olmalı), tutar '{metin(r['tutar'])}'. Satır atlandı."))
            continue
        tip_k = katla(r["tip"])
        degisken = tip_k.startswith(("degisken", "endeks", "tufe"))
        duyarsiz = tip_k.startswith(("faizsiz", "duyarsiz", "faize duyarsiz")) or tur_k.startswith("ozkaynak")
        vade, yf = tarih(r["vade"]), tarih(r["yf"])
        for a, ham, d in (("vade", r["vade"], vade), ("yeniden fiyatlama", r["yf"], yf)):
            if metin(ham) and not d:
                uy.append(uyari("Yüksek", "Tarih okunamadı", ad, f"Satır {r['_satir']}: {a} tarihi '{metin(ham)}'."))
        gun_t = None if duyarsiz else (yf or vade) if degisken else vade
        if not duyarsiz and gun_t is None:
            uy.append(uyari("Yüksek", "Yeniden fiyatlama tarihi yok", ad, f"Satır {r['_satir']}: faize duyarlı kalemde " +
                            ("yeniden fiyatlama veya vade" if degisken else "vade") + " tarihi yok. Kalem 'Faize duyarsız' sütununa alındı; vadesiz kalemler için "
                            "varsayımınıza göre bir tarih girin."))
            duyarsiz = True
        if degisken and not yf and vade:
            uy.append(uyari("Orta", "Değişken faizde yeniden fiyatlama tarihi yok", ad, "Vade tarihi kullanıldı; kalem olduğundan daha uzun vadeli görünebilir."))
        gun = None if duyarsiz else (gun_t - rapor).days
        if gun is not None and gun < 0:
            uy.append(uyari("Orta", "Tarihi geçmiş kalem", ad, f"Yeniden fiyatlama tarihi {gun_t:%d.%m.%Y} rapor tarihinden önce; 0-1 ay dilimine alındı."))
            gun = 0
        oran = para(r["oran"])
        kalemler.append({"ad": ad, "tur": tur, "tutar": t, "tip": "Duyarsız" if duyarsiz else "Değişken" if degisken else "Sabit", "gun": gun, "tarih": gun_t if not duyarsiz else None,
                         "dilim": DUYARSIZ if duyarsiz else dilim_bul(gun), "oran": oran, "pb": pb})
    return kalemler, uy


def analiz(kalemler: list[dict], soklar: list[int], ozkaynak: Decimal | None = None) -> dict:
    uy = []
    adlar = [d[0] for d in DILIMLER]
    tablo = {a: {"Varlık": Decimal(0), "Yükümlülük": Decimal(0)} for a in adlar + [DUYARSIZ]}
    for k in kalemler:
        tablo[k["dilim"]][k["tur"]] += k["tutar"]
    aktif = sum((k["tutar"] for k in kalemler if k["tur"] == "Varlık"), Decimal(0))
    pasif = sum((k["tutar"] for k in kalemler if k["tur"] == "Yükümlülük"), Decimal(0))
    kum, satirlar = Decimal(0), []
    for a in adlar:
        v, y = tablo[a]["Varlık"], tablo[a]["Yükümlülük"]
        kum += v - y
        satirlar.append({"dilim": a, "varlik": v, "yukumluluk": y, "gap": v - y, "kumulatif": kum, "gap_oran": yuvarla((v - y) / aktif * 100) if aktif else None,
                         "kum_oran": yuvarla(kum / aktif * 100) if aktif else None, "oran": yuvarla(v / y, "0.001") if y else None})
    if aktif != pasif:
        uy.append(uyari("Bilgi", "Bilanço denk değil", "Toplam", f"Varlıklar {bin_(aktif)}, yükümlülükler (özkaynak dahil) {bin_(pasif)}; fark {bin_(aktif - pasif)}. "
                                                                 "Dosya bilançonun tamamını içermiyorsa bu beklenir."))
    bir_yil = next(s for s in satirlar if s["dilim"] == "6-12 ay")
    sonuc = []
    for bp in soklar:
        dr = Decimal(bp) / 10000
        nfg, ed, dilim_nfg, dilim_ed = Decimal(0), Decimal(0), defaultdict(Decimal), defaultdict(Decimal)
        for k in kalemler:
            if k["gun"] is None:
                continue
            t = Decimal(k["gun"]) / 365
            isaret = 1 if k["tur"] == "Varlık" else -1
            if t < 1:
                x = isaret * k["tutar"] * dr * (1 - t)
                nfg += x
                dilim_nfg[k["dilim"]] += x
            y = -isaret * k["tutar"] * t / (1 + (k["oran"] or Decimal(0)) / 100) * dr
            ed += y
            dilim_ed[k["dilim"]] += y
        sonuc.append({"bp": bp, "nfg": yuvarla(nfg), "ed": yuvarla(ed), "ed_ozkaynak": yuvarla(ed / ozkaynak * 100) if ozkaynak else None,
                      "nfg_ozkaynak": yuvarla(nfg / ozkaynak * 100) if ozkaynak else None, "dilim_nfg": dict(dilim_nfg), "dilim_ed": dict(dilim_ed)})
    yon = "Varlığa duyarlı (faiz artışı net faiz gelirini artırır)" if bir_yil["kumulatif"] > 0 else \
          "Yükümlülüğe duyarlı (faiz artışı net faiz gelirini azaltır)" if bir_yil["kumulatif"] < 0 else "Dengeli"
    oransiz = [k["ad"] for k in kalemler if k["gun"] is not None and k["gun"] > 365 and k["oran"] is None]
    if oransiz:
        uy.append(uyari("Bilgi", "Faiz oranı girilmemiş uzun vadeli kalem", f"{len(oransiz)} kalem", "Ekonomik değer hesabında iskonto yapılmadı (etki bir miktar büyük çıkar): " +
                        ", ".join(oransiz[:8]) + (" …" if len(oransiz) > 8 else "")))
    return {"satirlar": satirlar, "duyarsiz": tablo[DUYARSIZ], "aktif": aktif, "pasif": pasif, "soklar": sonuc, "yon": yon, "kum_1y": bir_yil["kumulatif"],
            "uyarilar": uy, "kalemler": kalemler, "ozkaynak": ozkaynak}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

def rapor_yaz(cikti: Path, r: dict, rapor: date, pb: str) -> None:
    wb = Workbook()
    gp = wb.active
    gp.title = "Gap Tablosu"
    gp.append([f"Rapor tarihi {rapor:%d.%m.%Y}" + (f" · {pb}" if pb else "") + f" · 1 yıla kadar kümülatif gap: {bin_(r['kum_1y'])} → {r['yon']}"])
    gp["A1"].font = Font(bold=True)
    _baslik(gp, ["Yeniden Fiyatlama Dilimi", "Faize Duyarlı Varlık", "Faize Duyarlı Yükümlülük", "Gap", "Kümülatif Gap", "Gap ÷ Toplam Varlık %", "Kümülatif ÷ Toplam Varlık %",
                 "Varlık ÷ Yükümlülük"], (24, 20, 20, 20, 20, 13, 13, 12), "B3")
    for s in r["satirlar"]:
        _satir(gp, [s["dilim"], s["varlik"], s["yukumluluk"], s["gap"], s["kumulatif"], s["gap_oran"], s["kum_oran"], s["oran"]], para_s=(2, 3, 4, 5),
               dolgu={4: YESIL if s["gap"] > 0 else KIRMIZI if s["gap"] < 0 else None})
    dv, dy = r["duyarsiz"]["Varlık"], r["duyarsiz"]["Yükümlülük"]
    _satir(gp, [DUYARSIZ, dv, dy, "", "", "", "", ""], para_s=(2, 3))
    _satir(gp, ["Toplam", r["aktif"], r["pasif"], "", "", "", "", ""], para_s=(2, 3))
    for c in gp[gp.max_row]:
        c.font = Font(bold=True)
    gp.append([])
    gp.append(["Artı gap: o dilimde yeniden fiyatlanan varlık fazla (faiz artışı lehte). Eksi gap: yükümlülük fazla (faiz artışı aleyhte)."])

    sk = wb.create_sheet("Şok Etkisi")
    _baslik(sk, ["Paralel Şok (baz puan)", "Net Faiz Geliri Etkisi (12 ay)", "Ekonomik Değer Etkisi (yaklaşık)", "NFG Etkisi ÷ Özkaynak %", "ED Etkisi ÷ Özkaynak %"],
            (20, 26, 26, 18, 18))
    for s in r["soklar"]:
        _satir(sk, [s["bp"], s["nfg"], s["ed"], s["nfg_ozkaynak"], s["ed_ozkaynak"]], para_s=(2, 3),
               dolgu={2: KIRMIZI if s["nfg"] < 0 else YESIL if s["nfg"] > 0 else None, 3: KIRMIZI if s["ed"] < 0 else YESIL if s["ed"] > 0 else None})
    sk.append([])
    _baslik(sk, ["Dilim bazında", *[f"NFG {s['bp']:+d} bp" for s in r["soklar"]], *[f"ED {s['bp']:+d} bp" for s in r["soklar"]]], dondur=None)
    for a in [d[0] for d in DILIMLER]:
        _satir(sk, [a, *[yuvarla(s["dilim_nfg"].get(a, Decimal(0))) for s in r["soklar"]], *[yuvarla(s["dilim_ed"].get(a, Decimal(0))) for s in r["soklar"]]],
               para_s=tuple(range(2, 2 + 2 * len(r["soklar"]))))
    sk.append([])
    sk.append(["NFG: kalem yeniden fiyatlandığı günden yıl sonuna kadar yeni orandan faiz işler (bilanço sabit, anlık paralel şok). "
               "ED: tutar × süre ÷ (1 + faiz) × şok; ara ödemeler yok sayıldığı için uzun vadeli kalemlerde etki olduğundan büyük çıkar."])

    kl = wb.create_sheet("Kalemler")
    _baslik(kl, ["Kalem", "Tür", "Tutar", "Faiz Tipi", "Yeniden Fiyatlama Tarihi", "Gün", "Dilim", "Faiz Oranı %", "Para Birimi"], (40, 12, 20, 11, 14, 8, 16, 10, 9))
    for k in sorted(r["kalemler"], key=lambda k: (k["tur"], k["gun"] if k["gun"] is not None else 10**6)):
        _satir(kl, [k["ad"], k["tur"], k["tutar"], k["tip"], k["tarih"], k["gun"], k["dilim"], k["oran"], k["pb"]], para_s=(3,), tarih_s=(5,))
    kl.auto_filter.ref = f"A1:I{kl.max_row}"
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, rapor: date, soklar: list[int], ozkaynak: Decimal | None = None, pb: str = "") -> dict:
    kalemler, uy = oku(yol, rapor, pb)
    if not kalemler:
        raise ValueError(f"{yol.name}: okunabilir kalem yok" + (f" ({pb} için)" if pb else ""))
    birimler = sorted({k["pb"] for k in kalemler})
    r = analiz(kalemler, soklar, ozkaynak)
    if len(birimler) > 1:
        r["uyarilar"].append(uyari("Orta", "Birden çok para birimi", ", ".join(birimler), "Tutarlar aynı para birimine çevrilmiş kabul edilip toplandı. Faiz şokları para birimine "
                                                                                       "göre farklıysa --para-birimi ile ayrı ayrı çalıştırın."))
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, rapor, pb)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Faize duyarlı kalemleri yeniden fiyatlama vadesine göre dağıtır; gap tablosu ve faiz şoku etkisini hesaplar.")
    p.add_argument("--kalemler", type=Path, default=ORNEK / "kalemler.csv",
                   help="Kalem, Tür (Varlık / Yükümlülük), Tutar, Faiz Tipi (Sabit / Değişken / Faizsiz), Vade Tarihi, Yeniden Fiyatlama Tarihi, Faiz Oranı, Para Birimi")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 30.09.2026)")
    p.add_argument("--sok", default="100,-100", help="Paralel şoklar, baz puan, virgülle (varsayılan 100,-100)")
    p.add_argument("--ozkaynak", help="Özkaynak tutarı (etkilerin özkaynağa oranı için)")
    p.add_argument("--para-birimi", default="", help="Yalnız bu para birimindeki kalemler (ör. TL)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "faiz_gap_analizi.xlsx")
    a = p.parse_args(argv)
    if not a.kalemler.exists():
        print(f"[X] Dosya bulunamadı: {a.kalemler}")
        return 1
    ornek = a.kalemler == ORNEK / "kalemler.csv"
    rapor = tarih(a.tarih) if a.tarih else (date(2026, 9, 30) if ornek else date.today())
    oz = para(a.ozkaynak) if a.ozkaynak else (Decimal("2400000000") if ornek else None)
    try:
        soklar = [int(x) for x in a.sok.replace(" ", "").split(",") if x]
        if not rapor or not soklar or (a.ozkaynak and (oz is None or oz <= 0)):
            raise ValueError
    except ValueError:
        print("[X] --tarih GG.AA.YYYY; --sok baz puan listesi (ör. 100,-100); --ozkaynak pozitif sayı olmalı")
        return 2
    try:
        r = calistir(a.kalemler, a.cikti, rapor, soklar, oz, a.para_birimi.upper().replace("TRY", "TL"))
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    b = bin_
    print(f"[OK] {len(r['kalemler'])} kalem · 1 yıla kadar kümülatif gap {b(r['kum_1y'])} · {r['yon']}")
    for s in r["soklar"]:
        print(f"     {s['bp']:+d} bp: net faiz geliri {b(s['nfg'])} · ekonomik değer {b(s['ed'])}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['kim']} · {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
