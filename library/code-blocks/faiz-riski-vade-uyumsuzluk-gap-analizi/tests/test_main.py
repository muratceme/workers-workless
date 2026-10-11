import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
RAPOR = date(2026, 9, 30)


def kalem(ad, tur, tutar, gun, oran=None):
    return {"ad": ad, "tur": tur, "tutar": D(str(tutar)), "tip": "Sabit" if gun is not None else "Duyarsız", "gun": gun, "tarih": None,
            "dilim": main.DUYARSIZ if gun is None else main.dilim_bul(gun), "oran": D(str(oran)) if oran is not None else None, "pb": "TL"}


class HesapTesti(unittest.TestCase):
    def test_dilim_sinirlari(self):
        self.assertEqual([main.dilim_bul(g) for g in (0, 30, 31, 91, 92, 182, 183, 365, 366, 730, 731, 1826, 1827, 9999)],
                         ["0-1 ay", "0-1 ay", "1-3 ay", "1-3 ay", "3-6 ay", "3-6 ay", "6-12 ay", "6-12 ay", "1-2 yıl", "1-2 yıl", "2-5 yıl", "2-5 yıl", "5 yıl+", "5 yıl+"])

    def test_gap_tablosu(self):
        k = [kalem("A", "Varlık", 500, 10), kalem("B", "Yükümlülük", 800, 20), kalem("C", "Varlık", 600, 200), kalem("D", "Yükümlülük", 100, 400),
             kalem("Nakit", "Varlık", 100, None), kalem("Özkaynak", "Yükümlülük", 300, None)]
        r = main.analiz(k, [100])
        t = {s["dilim"]: (s["gap"], s["kumulatif"]) for s in r["satirlar"]}
        self.assertEqual((t["0-1 ay"], t["6-12 ay"], t["1-2 yıl"], t["5 yıl+"]), ((D(-300), D(-300)), (D(600), D(300)), (D(-100), D(200)), (D(0), D(200))))
        self.assertEqual((r["aktif"], r["pasif"], r["duyarsiz"]), (D(1200), D(1200), {"Varlık": D(100), "Yükümlülük": D(300)}))
        self.assertEqual(r["satirlar"][0]["gap_oran"], D("-25.00"))                  # −300 / 1.200
        self.assertTrue(r["yon"].startswith("Varlığa duyarlı"))                       # 1 yıla kadar kümülatif +300
        self.assertEqual([u["tur"] for u in r["uyarilar"]], ["Faiz oranı girilmemiş uzun vadeli kalem"])   # D: 400 gün, oran yok; bilanço denk

    def test_net_faiz_geliri(self):
        r = main.analiz([kalem("V", "Varlık", 1000, 73), kalem("Y", "Yükümlülük", 2000, 0), kalem("Uzun", "Varlık", 5000, 400, 20)], [100, -200])
        self.assertEqual(r["soklar"][0]["nfg"], D("-12.00"))                          # +1000 × %1 × (1 − 0,2) − 2000 × %1 × 1
        self.assertEqual(r["soklar"][1]["nfg"], D("24.00"))
        self.assertEqual(r["soklar"][0]["dilim_nfg"], {"1-3 ay": D("8.000"), "0-1 ay": D("-20")})
        self.assertTrue(r["yon"].startswith("Yükümlülüğe duyarlı"))

    def test_ekonomik_deger(self):
        r = main.analiz([kalem("Tahvil", "Varlık", 1000, 730, 25), kalem("Borç", "Yükümlülük", 400, 365)], [100], D(100))
        self.assertEqual(r["soklar"][0]["ed"], D("-12.00"))                           # −1000 × 2 / 1,25 × %1 + 400 × 1 × %1
        self.assertEqual((r["soklar"][0]["ed_ozkaynak"], r["soklar"][0]["nfg_ozkaynak"]), (D("-12.00"), D("0.00")))
        self.assertEqual([u["tur"] for u in r["uyarilar"]], ["Bilanço denk değil"])
        r = main.analiz([kalem("Uzun", "Varlık", 1000, 730)], [100])
        self.assertEqual(r["soklar"][0]["ed"], D("-20.00"))                           # oran yoksa iskonto yapılmaz
        self.assertIn("Faiz oranı girilmemiş uzun vadeli kalem", [u["tur"] for u in r["uyarilar"]])


class DosyaTesti(unittest.TestCase):
    def test_okuma_kurallari(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "k.csv"
            y.write_text("Kalem;Tür;Tutar;Faiz Tipi;Vade Tarihi;Yeniden Fiyatlama Tarihi;Faiz Oranı;Para Birimi\n"
                         "Sabit;Varlık;1.000;Sabit;30.10.2026;;40;TL\n"
                         "Değişken;Aktif;2.000;Değişken;30.09.2029;31.12.2026;;TRY\n"
                         "Değişken YF yok;Varlık;500;Değişken;30.09.2027;;;TL\n"
                         "Geçmiş;Pasif;700;Sabit;15.09.2026;;;TL\n"
                         "Vadesiz;Yükümlülük;900;Sabit;;;;TL\n"
                         "Özkaynak;Özkaynak;300;;;;;TL\n"
                         "Dolar;Varlık;400;Sabit;30.11.2026;;;USD\n"
                         "Bozuk;Gelir;100;Sabit;;;;TL\n", encoding="utf-8")
            k, uy = main.oku(y, RAPOR)
            self.assertEqual([(x["ad"], x["tur"], x["tip"], x["gun"], x["dilim"]) for x in k],
                             [("Sabit", "Varlık", "Sabit", 30, "0-1 ay"), ("Değişken", "Varlık", "Değişken", 92, "3-6 ay"), ("Değişken YF yok", "Varlık", "Değişken", 365, "6-12 ay"),
                              ("Geçmiş", "Yükümlülük", "Sabit", 0, "0-1 ay"), ("Vadesiz", "Yükümlülük", "Duyarsız", None, "Faize duyarsız"),
                              ("Özkaynak", "Yükümlülük", "Duyarsız", None, "Faize duyarsız"), ("Dolar", "Varlık", "Sabit", 61, "1-3 ay")])
            self.assertEqual(sorted(u["tur"] for u in uy), ["Değişken faizde yeniden fiyatlama tarihi yok", "Satır okunamadı", "Tarihi geçmiş kalem", "Yeniden fiyatlama tarihi yok"])
            self.assertEqual([x["ad"] for x in main.oku(y, RAPOR, "USD")[0]], ["Dolar"])
            arg = ["--kalemler", str(y), "--tarih", "30.09.2026", "--cikti", str(Path(tmp) / "x.xlsx")]
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(arg + ["--para-birimi", "EUR"]), 1)
            self.assertEqual(main.main(arg + ["--sok", "yüz"]), 2)
            self.assertEqual(main.main(["--kalemler", str(Path(tmp) / "yok.csv")]), 1)

    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "g.xlsx"
            r = main.calistir(ORNEK / "kalemler.csv", cikti, RAPOR, [100, -100], D("2400000000"))
            self.assertEqual((len(r["kalemler"]), r["aktif"], r["pasif"], r["kum_1y"]), (26, D("31500000000"), D("31500000000"), D("-7300000000")))
            self.assertEqual(r["satirlar"][-1]["kumulatif"], D("5100000000"))         # duyarlı varlık − duyarlı yükümlülük = 28,6 − 23,5 milyar
            self.assertEqual(r["soklar"][0]["nfg"], -r["soklar"][1]["nfg"])
            self.assertTrue(r["soklar"][0]["nfg"] < 0 and r["soklar"][0]["ed"] < 0)
            self.assertEqual(r["uyarilar"], [])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Gap Tablosu", "Şok Etkisi", "Kalemler", "Uyarılar"])
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "o.xlsx")]), 0)


if __name__ == "__main__":
    unittest.main()
