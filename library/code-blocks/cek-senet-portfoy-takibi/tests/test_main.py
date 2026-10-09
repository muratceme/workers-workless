import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 9)


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "c.xlsx"
        cls.s = main.calistir(ORNEK / "portfoy.csv", cls.cikti, BUGUN, ORNEK / "banka_bakiyeleri.csv")
        cls.u = {(u["kim"], u["tur"]) for u in cls.s["uyarilar"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_portfoy_toplamlari(self):
        al = sum(e.tutar for e in self.s["acik_alinan"] if e.doviz == "TRY")
        ve = sum(e.tutar for e in self.s["acik_verilen"])
        self.assertEqual((al, ve), (D("819500.00"), D("567050.00")))       # ciro, tahsil, karşılıksız, ödendi hariç
        self.assertEqual(round(self.s["ort_vade"]["TRY"]), 25)

    def test_uyarilar(self):
        self.assertTrue({("D-330981", "İbraz süresi"), ("S-2026-014", "Vadesi geçmiş"), ("E-901244", "Karşılıksız"), ("E-901245", "Riskli keşideci"),
                         ("X Bankası", "Karşılık açığı"), ("Örnek Yapı Ltd.", "Yoğunlaşma"), ("H-612001", "Ciro riski")} <= self.u)
        self.assertNotIn(("Y Bankası", "Karşılık açığı"), self.u)        # 240.000 + tahsile verilen 84.000 ≥ 176.300

    def test_haftalik_ve_karsilik(self):
        h = self.s["haftalar"][1]["doviz"]["TRY"]                        # 12–18 Ekim
        self.assertEqual((h["tahsilat"], h["odeme"]), (D("150000.00"), D("318750.00")))
        self.assertEqual(self.s["haftalar"][2]["doviz"]["TRY"]["teminat"], D("95000.00"))
        self.assertEqual(self.s["gecmis"]["TRY"]["tahsilat"], D("77000.00"))
        k = [x["acik"] for x in self.s["karsilik"] if x["banka"] == "X Bankası"]
        self.assertEqual(k, [D("13750.00"), D("85750.00")])

    def test_excel(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Özet", "Haftalık Vade", "Banka Karşılık", "Portföy", "Keşideci Riski", "Uyarılar"])
        self.assertEqual(wb["Portföy"].max_row, 19)


class KuralTesti(unittest.TestCase):
    def test_ibraz_suresi_ve_verilen(self):
        with tempfile.TemporaryDirectory() as tmp:
            s = main.calistir(ORNEK / "portfoy.csv", Path(tmp) / "o.xlsx", date(2026, 10, 21), ibraz_gun=30)
            u = {(x["kim"], x["tur"]) for x in s["uyarilar"]}
            self.assertIn(("D-330981", "Vadesi geçmiş"), u)              # 30 gün içinde: ibraz süresi uyarısı yok
            self.assertIn(("K-000451", "Ödenmemiş görünüyor"), u)
            self.assertIn(("B-552018", "Sonuç bekleniyor"), u)
            self.assertIn(("X Bankası", "Bakiye yok"), u)
            self.assertEqual(main.pazartesi(date(2026, 10, 9)), date(2026, 10, 5))

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--bugun", "dün"]), 2)
            self.assertEqual(main.main(["--portfoy", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
