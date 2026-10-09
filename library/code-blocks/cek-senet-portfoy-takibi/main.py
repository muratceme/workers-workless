"""
Çek-Senet Portföy Takibi — Workers / Workless kod bloğu
Finans › Finans Uzmanı

Alınan ve verilen çek-senetleri vade, banka, keşideci ve durum bazında izler; haftalık vade listesi çıkarır:
  - Açık kayıtlar: alınan → Portföyde, Tahsile Verildi, Teminata Verildi; verilen → Verildi (ödenmemiş).
    Kapalı: Tahsil Edildi, Ödendi, İade. Ciro edilen çek portföyden çıkar ama vadesine kadar müracaat riski izlenir.
  - Haftalık vade listesi (Pazartesi başlangıçlı, --hafta): beklenen tahsilat (teminattakiler ayrı), ödenecek, net.
  - Banka karşılık kontrolü: verilen çeklerin vadeye göre kümülatif tutarı, banka bakiyesi + o bankaya tahsile
    verilmiş ve o tarihe kadar vadesi gelen alınan çeklerle karşılaştırılır; açık varsa ilk eksik tarihi verilir.
  - Kontroller: vadesi + ibraz süresi (varsayılan 10 gün, TTK md. 796) geçmiş hâlâ portföydeki alınan çek, vadesi
    geçmiş açık senet / verilen çek, karşılıksız / protestolu kayıt ve aynı keşidecinin diğer açık evrakı, keşideci
    yoğunlaşması (açık alınan portföyün %30'u), ciro edilen çeklerde süren müracaat riski.
  - Ağırlıklı ortalama vade (gün) ve döviz bazında toplamlar.
Rapor: özet, haftalık vade listesi, banka karşılık, açık portföy, keşideci riski, uyarılar. İnternete bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: 18 evrak, bugün 09.10.2026
    python main.py --portfoy portfoy.xlsx --banka-bakiyeleri bakiyeler.xlsx --hafta 12
    python main.py --portfoy portfoy.xlsx --bugun 01.11.2026 --ibraz-gun 30
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
K2 = Decimal("0.01")
SIFIR = Decimal(0)
ACIK_ALINAN = {"portfoyde": "Portföyde", "tahsile verildi": "Tahsile Verildi", "teminata verildi": "Teminata Verildi"}
SORUNLU = {"karsiliksiz": "Karşılıksız", "protestolu": "Protestolu", "protesto": "Protestolu", "takipte": "Takipte"}

SUTUNLAR = {"tur": ("tur", "evrak turu"), "no": ("seri no", "cek no", "senet no", "no"), "kesideci": ("kesideci borclu", "kesideci", "borclu"),
            "cari": ("cari", "cari unvan", "musteri tedarikci"), "banka": ("banka",), "vade": ("vade", "vade tarihi"), "tutar": ("tutar",),
            "doviz": ("doviz", "para birimi"), "durum": ("durum",), "durum_tarihi": ("durum tarihi",)}
BANKA_SUTUNLARI = {"banka": ("banka",), "bakiye": ("bakiye", "kullanilabilir bakiye"), "doviz": ("doviz", "para birimi")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tl(x: Decimal | None) -> str:
    return "—" if x is None else f"{x.quantize(K2, ROUND_HALF_UP):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Evrak:
    yon: str             # alinan | verilen
    cins: str            # cek | senet
    no: str
    kesideci: str
    cari: str
    banka: str
    vade: date
    tutar: Decimal
    doviz: str
    durum: str
    durum_tarihi: date | None

    @property
    def durum_k(self) -> str:
        return katla(self.durum)

    @property
    def acik(self) -> bool:
        if self.yon == "alinan":
            return self.durum_k in ACIK_ALINAN
        return self.durum_k.startswith(("verildi", "portfoyde", "kesildi", "odenecek")) or self.durum_k == ""

    @property
    def sorunlu(self) -> bool:
        return any(self.durum_k.startswith(k) for k in SORUNLU)

    @property
    def ciro(self) -> bool:
        return self.durum_k.startswith("ciro")

    @property
    def tur_ad(self) -> str:
        return f"{'Alınan' if self.yon == 'alinan' else 'Verilen'} {'Çek' if self.cins == 'cek' else 'Senet'}"


def oku(yol: Path) -> tuple[list[Evrak], list[dict]]:
    sonuc, uy = [], []
    for r in kayitlar(yol, SUTUNLAR, ("tur", "vade", "tutar")):
        t = katla(r.get("tur"))
        v, m = tarih(r.get("vade")), para(r.get("tutar"))
        yon = "alinan" if t.startswith(("alinan", "musteri")) else "verilen" if t.startswith(("verilen", "kendi", "borc")) else None
        cins = "senet" if "senet" in t or "bono" in t else "cek"
        if not yon or not v or m is None:
            if any(r.get(k) for k in ("no", "tutar")):
                uy.append({"onem": "Orta", "tur": "Okunamayan satır", "kim": metin(r.get("no")), "aciklama": f"Satır {r['_satir']}: tür, vade veya tutar okunamadı"})
            continue
        sonuc.append(Evrak(yon, cins, metin(r.get("no")), metin(r.get("kesideci")), metin(r.get("cari")), metin(r.get("banka")), v, m,
                           (metin(r.get("doviz")) or "TRY").upper().replace("TL", "TRY"), metin(r.get("durum")), tarih(r.get("durum_tarihi"))))
    return sonuc, uy


# ----------------------------------------------------------------------------
# Analiz
# ----------------------------------------------------------------------------

def pazartesi(d: date) -> date:
    return d - timedelta(days=d.weekday())


def analiz_et(evraklar: list[Evrak], bugun: date, hafta: int, ibraz_gun: int, bakiyeler: dict[tuple[str, str], Decimal], yogunluk: Decimal) -> dict:
    uy = []

    def uyar(onem, tur, e: Evrak | None, aciklama, kim=""):
        uy.append({"onem": onem, "tur": tur, "kim": e.no if e else kim, "aciklama": aciklama})

    acik_alinan = [e for e in evraklar if e.yon == "alinan" and e.acik]
    acik_verilen = [e for e in evraklar if e.yon == "verilen" and e.acik]
    sorunlu = [e for e in evraklar if e.sorunlu]
    for e in acik_alinan:
        if e.vade < bugun:
            if e.cins == "cek" and e.durum_k == "portfoyde" and e.vade + timedelta(days=ibraz_gun) < bugun:
                uyar("Yüksek", "İbraz süresi", e, f"Vadesi {e.vade:%d.%m.%Y}, hâlâ portföyde; {ibraz_gun} günlük ibraz süresi geçmiş olabilir (TTK md. 796). "
                     "Süresinde ibraz edilmeyen çekte cirantalara başvuru hakkı kaybedilebilir")
            elif e.durum_k == "portfoyde":
                uyar("Yüksek", "Vadesi geçmiş", e, f"{e.tur_ad} vadesi {e.vade:%d.%m.%Y}, hâlâ portföyde; tahsile verin veya durumunu güncelleyin")
            else:
                uyar("Orta", "Sonuç bekleniyor", e, f"{e.durum} durumunda, vadesi {e.vade:%d.%m.%Y} geçti; bankadan sonucu teyit edin")
    for e in acik_verilen:
        if e.vade < bugun:
            uyar("Yüksek", "Ödenmemiş görünüyor", e, f"Verilen {e.cins} vadesi {e.vade:%d.%m.%Y} geçti, durum '{e.durum}'; ödendiyse kaydı kapatın")
    for e in sorunlu:
        uyar("Yüksek", e.durum, e, f"{e.kesideci} · {tl(e.tutar)} {e.doviz}; hukuki takip ve cari risk değerlendirmesi")
        for d in acik_alinan:
            if katla(d.kesideci) == katla(e.kesideci):
                uyar("Yüksek", "Riskli keşideci", d, f"Aynı keşidecinin ({e.kesideci}) {e.no} no'lu evrakı {e.durum.lower()}; bu evrakın tahsil riski yüksek")
    for e in evraklar:
        if e.ciro and e.vade >= bugun:
            uyar("Bilgi", "Ciro riski", e, f"{e.cari or 'Üçüncü kişiye'} ciro edildi; vadesi {e.vade:%d.%m.%Y}. Karşılıksız çıkarsa müracaat yoluyla geri gelebilir")
    # Haftalık vade listesi
    bas = pazartesi(bugun)
    haftalar = []
    for i in range(hafta):
        h1, h2 = bas + timedelta(days=7 * i), bas + timedelta(days=7 * i + 6)
        g = defaultdict(lambda: defaultdict(lambda: SIFIR))
        for e in acik_alinan:
            if h1 <= e.vade <= h2:
                g[e.doviz]["teminat" if e.durum_k == "teminata verildi" else "tahsilat"] += e.tutar
        for e in acik_verilen:
            if h1 <= e.vade <= h2:
                g[e.doviz]["odeme"] += e.tutar
        haftalar.append({"bas": h1, "bit": h2, "doviz": g})
    gecmis = defaultdict(lambda: defaultdict(lambda: SIFIR))
    for e in acik_alinan + acik_verilen:
        if e.vade < bas:
            gecmis[e.doviz]["odeme" if e.yon == "verilen" else "tahsilat"] += e.tutar
    # Banka karşılık
    karsilik = []
    for (banka, dv), bakiye in sorted(bakiyeler.items()):
        cekler = sorted([e for e in acik_verilen if e.cins == "cek" and katla(e.banka) == katla(banka) and e.doviz == dv], key=lambda e: e.vade)
        if not cekler:
            continue
        kum, ilk_acik = SIFIR, None
        for e in cekler:
            kum += e.tutar
            gelen = sum((a.tutar for a in acik_alinan if a.durum_k == "tahsile verildi" and katla(a.banka) == katla(banka) and a.doviz == dv
                         and a.vade <= e.vade), SIFIR)
            acik = kum - bakiye - gelen
            karsilik.append({"banka": banka, "doviz": dv, "e": e, "kum": kum, "bakiye": bakiye, "gelen": gelen, "acik": max(acik, SIFIR)})
            if acik > 0 and ilk_acik is None:
                ilk_acik = (e, acik)
        if ilk_acik:
            e, acik = ilk_acik
            uy.append({"onem": "Yüksek", "tur": "Karşılık açığı", "kim": banka, "aciklama": f"{e.vade:%d.%m.%Y} vadeli {e.no} çeki için {tl(acik)} {dv} "
                       "karşılık eksik (bakiye + tahsile verilen çekler yetmiyor). Karşılıksız çek, 5941 sayılı Çek Kanunu kapsamında "
                       "yaptırımlara yol açabilir"})
    tanimli = {(katla(k[0]), k[1]) for k in bakiyeler}
    eksik = {}
    for e in acik_verilen:
        if e.cins == "cek" and e.banka and (katla(e.banka), e.doviz) not in tanimli:
            eksik.setdefault((katla(e.banka), e.doviz), e.banka)
    for (_, dv), ad in sorted(eksik.items()):
        uy.append({"onem": "Bilgi", "tur": "Bakiye yok", "kim": ad, "aciklama": f"Bu bankadan verilen {dv} çekler var ama banka bakiyesi verilmedi"})
    # Keşideci yoğunlaşması ve ortalama vade
    risk = []
    for dv in sorted({e.doviz for e in acik_alinan}):
        lst = [e for e in acik_alinan if e.doviz == dv]
        top = sum((e.tutar for e in lst), SIFIR)
        g = defaultdict(list)
        for e in lst:
            g[katla(e.kesideci)].append(e)
        for k, es in sorted(g.items(), key=lambda i: -sum(e.tutar for e in i[1])):
            t = sum((e.tutar for e in es), SIFIR)
            pay = t / top if top else SIFIR
            sorun = any(katla(s.kesideci) == k for s in sorunlu)
            risk.append({"kesideci": es[0].kesideci, "doviz": dv, "adet": len(es), "tutar": t, "pay": pay, "en_yakin": min(e.vade for e in es), "sorunlu": sorun})
            if pay > yogunluk and len(lst) > 2:
                uy.append({"onem": "Orta", "tur": "Yoğunlaşma", "kim": es[0].kesideci, "aciklama": f"Açık alınan {dv} portföyün %{pay * 100:.0f}'i tek keşidecide "
                           f"({tl(t)} {dv}, {len(es)} evrak)"})
    ort_vade = {}
    for dv in sorted({e.doviz for e in acik_alinan}):
        lst = [e for e in acik_alinan if e.doviz == dv]
        top = sum((e.tutar for e in lst), SIFIR)
        ort_vade[dv] = (sum((e.tutar * (e.vade - bugun).days for e in lst), SIFIR) / top) if top else None
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    uy.sort(key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))
    return {"evraklar": evraklar, "acik_alinan": acik_alinan, "acik_verilen": acik_verilen, "haftalar": haftalar, "gecmis": gecmis,
            "karsilik": karsilik, "risk": risk, "ort_vade": ort_vade, "uyarilar": uy, "bugun": bugun, "sorunlu": sorunlu}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
UST = Alignment(vertical="top", wrap_text=True)
PF = "#,##0.00"


def _baslik(ws, basliklar, genislik=()):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def rapor_yaz(cikti: Path, s: dict) -> None:
    wb = Workbook()
    oz = wb.active
    oz.title = "Özet"
    _baslik(oz, ["Gösterge", "Döviz", "Adet", "Tutar", "Not"], (36, 7, 7, 16, 60))
    for dv in sorted({e.doviz for e in s["evraklar"]}):
        for ad, lst, n in (("Açık alınan çek / senet", [e for e in s["acik_alinan"] if e.doviz == dv], f"Ağırlıklı ortalama vade: "
                            + (f"{s['ort_vade'][dv]:.0f} gün" if s["ort_vade"].get(dv) is not None else "—")),
                           ("  · Portföyde", [e for e in s["acik_alinan"] if e.doviz == dv and e.durum_k == "portfoyde"], ""),
                           ("  · Tahsile verildi", [e for e in s["acik_alinan"] if e.doviz == dv and e.durum_k == "tahsile verildi"], ""),
                           ("  · Teminata verildi", [e for e in s["acik_alinan"] if e.doviz == dv and e.durum_k == "teminata verildi"], "Nakit girişi teminat hesabına"),
                           ("Açık verilen çek / senet", [e for e in s["acik_verilen"] if e.doviz == dv], ""),
                           ("Karşılıksız / protestolu", [e for e in s["sorunlu"] if e.doviz == dv], "Takip")):
            if lst or not ad.startswith("  "):
                oz.append([ad, dv, len(lst), float(sum((e.tutar for e in lst), SIFIR)), n])
                oz.cell(oz.max_row, 4).number_format = PF
    oz.append([])
    oz.append([f"Rapor tarihi: {s['bugun']:%d.%m.%Y}"])

    hv = wb.create_sheet("Haftalık Vade")
    dvler = sorted({d for h in s["haftalar"] for d in h["doviz"]} | set(s["gecmis"]))
    bas = ["Hafta"]
    for dv in dvler:
        bas += [f"Tahsilat ({dv})", f"Teminatta ({dv})", f"Ödeme ({dv})", f"Net ({dv})", f"Kümülatif Net ({dv})"]
    _baslik(hv, bas, [24] + [14] * 5 * len(dvler))
    kum = defaultdict(lambda: SIFIR)
    gec = ["Vadesi geçmiş, açık"]
    for dv in dvler:
        g = s["gecmis"].get(dv, {})
        net = g.get("tahsilat", SIFIR) - g.get("odeme", SIFIR)
        gec += [float(g.get("tahsilat", SIFIR)), None, float(g.get("odeme", SIFIR)), float(net), None]
    hv.append(gec)
    for h in s["haftalar"]:
        satir = [f"{h['bas']:%d.%m} – {h['bit']:%d.%m.%Y}"]
        for dv in dvler:
            g = h["doviz"].get(dv, {})
            net = g.get("tahsilat", SIFIR) - g.get("odeme", SIFIR)
            kum[dv] += net
            satir += [float(g.get("tahsilat", SIFIR)), float(g.get("teminat", SIFIR)), float(g.get("odeme", SIFIR)), float(net), float(kum[dv])]
        hv.append(satir)
    for row in hv.iter_rows(min_row=2, min_col=2):
        for c in row:
            c.number_format = PF
            if c.value is not None and c.value < 0 and (c.column - 2) % 5 >= 3:
                c.fill = PatternFill("solid", fgColor="FDE2E1")
    if "TRY" in dvler:
        j = 2 + dvler.index("TRY") * 5
        gr = BarChart()
        gr.title, gr.height, gr.width = "Haftalık tahsilat / ödeme (TRY)", 8, 20
        gr.add_data(Reference(hv, min_col=j, max_col=j, min_row=1, max_row=hv.max_row), titles_from_data=True)
        gr.add_data(Reference(hv, min_col=j + 2, max_col=j + 2, min_row=1, max_row=hv.max_row), titles_from_data=True)
        gr.set_categories(Reference(hv, min_col=1, min_row=2, max_row=hv.max_row))
        hv.add_chart(gr, f"A{hv.max_row + 3}")

    bk = wb.create_sheet("Banka Karşılık")
    _baslik(bk, ["Banka", "Döviz", "Çek No", "Vade", "Tutar", "Kümülatif Çek", "Banka Bakiyesi", "Tahsile Verilen (vadesi gelen)", "Karşılık Açığı"],
            (14, 7, 12, 11, 14, 15, 15, 16, 14))
    for x in s["karsilik"]:
        bk.append([x["banka"], x["doviz"], x["e"].no, x["e"].vade, float(x["e"].tutar), float(x["kum"]), float(x["bakiye"]), float(x["gelen"]), float(x["acik"])])
        bk.cell(bk.max_row, 4).number_format = "DD.MM.YYYY"
        for j in range(5, 10):
            bk.cell(bk.max_row, j).number_format = PF
        if x["acik"]:
            bk.cell(bk.max_row, 9).fill = PatternFill("solid", fgColor="FDE2E1")

    ap = wb.create_sheet("Portföy")
    _baslik(ap, ["Tür", "Seri No", "Keşideci / Borçlu", "Cari", "Banka", "Vade", "Kalan Gün", "Tutar", "Döviz", "Durum", "Durum Tarihi", "Açık", "Aksiyon"],
            (13, 12, 26, 26, 12, 11, 9, 14, 6, 16, 11, 6, 22))
    for e in sorted(s["evraklar"], key=lambda e: (not e.acik, e.vade)):
        ap.append([e.tur_ad, e.no, e.kesideci, e.cari, e.banka, e.vade, (e.vade - s["bugun"]).days, float(e.tutar), e.doviz, e.durum, e.durum_tarihi,
                   "Evet" if e.acik else "", ""])
        r = ap.max_row
        ap.cell(r, 6).number_format = ap.cell(r, 11).number_format = "DD.MM.YYYY"
        ap.cell(r, 8).number_format = PF
        if e.sorunlu or (e.acik and e.vade < s["bugun"]):
            ap.cell(r, 7).fill = PatternFill("solid", fgColor="FDE2E1")
        ap.cell(r, 13).fill = PatternFill("solid", fgColor="FFF4CE")
    ap.auto_filter.ref = f"A1:M{ap.max_row}"

    kr = wb.create_sheet("Keşideci Riski")
    _baslik(kr, ["Keşideci", "Döviz", "Açık Evrak", "Tutar", "Portföy Payı", "En Yakın Vade", "Karşılıksız / Protestolu Geçmişi"], (28, 7, 9, 15, 11, 12, 16))
    for x in s["risk"]:
        kr.append([x["kesideci"], x["doviz"], x["adet"], float(x["tutar"]), float(x["pay"]), x["en_yakin"], "Evet" if x["sorunlu"] else ""])
        kr.cell(kr.max_row, 4).number_format = PF
        kr.cell(kr.max_row, 5).number_format = "0.0%"
        kr.cell(kr.max_row, 6).number_format = "DD.MM.YYYY"
        if x["sorunlu"]:
            kr.cell(kr.max_row, 7).fill = PatternFill("solid", fgColor="FDE2E1")

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Evrak / Kim", "Açıklama", "Aksiyon"], (9, 20, 16, 90, 22))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["kim"], u["aciklama"], ""])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).alignment = UST
        uy.cell(uy.max_row, 5).fill = PatternFill("solid", fgColor="FFF4CE")
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(portfoy: Path, cikti: Path, bugun: date, banka_yolu: Path | None = None, hafta: int = 8, ibraz_gun: int = 10,
             yogunluk: Decimal = Decimal("0.30")) -> dict:
    evraklar, uy = oku(portfoy)
    if not evraklar:
        raise ValueError(f"{portfoy.name}: evrak bulunamadı")
    bakiyeler = {}
    if banka_yolu:
        for r in kayitlar(banka_yolu, BANKA_SUTUNLARI, ("banka", "bakiye")):
            b = para(r.get("bakiye"))
            if metin(r.get("banka")) and b is not None:
                bakiyeler[(metin(r["banka"]), (metin(r.get("doviz")) or "TRY").upper().replace("TL", "TRY"))] = b
    s = analiz_et(evraklar, bugun, hafta, ibraz_gun, bakiyeler, yogunluk)
    s["uyarilar"] = uy + s["uyarilar"]
    rapor_yaz(cikti, s)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Alınan ve verilen çek-senetleri vade, banka ve durum bazında izler; haftalık vade listesi çıkarır.")
    p.add_argument("--portfoy", type=Path, default=ORNEK / "portfoy.csv",
                   help="Tür (Alınan/Verilen Çek/Senet), Seri No, Keşideci/Borçlu, Cari, Banka, Vade, Tutar, Döviz, Durum, Durum Tarihi")
    p.add_argument("--banka-bakiyeleri", type=Path, help="İsteğe bağlı: Banka, Bakiye, Döviz (verilen çek karşılık kontrolü)")
    p.add_argument("--bugun", help="Rapor tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 09.10.2026)")
    p.add_argument("--hafta", type=int, default=8, help="Vade listesi hafta sayısı (varsayılan 8)")
    p.add_argument("--ibraz-gun", type=int, default=10, help="Alınan çek ibraz süresi, gün (varsayılan 10; başka yerde ödenecek çekte 30)")
    p.add_argument("--yogunluk", type=float, default=30, help="Tek keşideci payı bu %%'yi aşarsa uyarı (varsayılan 30)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "cek_senet_portfoyu.xlsx")
    a = p.parse_args(argv)
    ornek = a.portfoy == ORNEK / "portfoy.csv"
    banka = a.banka_bakiyeleri or (ORNEK / "banka_bakiyeleri.csv" if ornek else None)
    for y in (a.portfoy, banka):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    bugun = tarih(a.bugun) if a.bugun else (date(2026, 10, 9) if ornek else date.today())
    if not bugun:
        print("[X] --bugun GG.AA.YYYY biçiminde olmalı")
        return 2
    try:
        s = calistir(a.portfoy, a.cikti, bugun, banka, a.hafta, a.ibraz_gun, Decimal(str(a.yogunluk)) / 100)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    for dv in sorted({e.doviz for e in s["evraklar"]}):
        al = sum((e.tutar for e in s["acik_alinan"] if e.doviz == dv), SIFIR)
        ve = sum((e.tutar for e in s["acik_verilen"] if e.doviz == dv), SIFIR)
        print(f"[OK] {dv}: açık alınan {tl(al)} · açık verilen {tl(ve)}" + (f" · ortalama vade {s['ort_vade'][dv]:.0f} gün" if s["ort_vade"].get(dv) else ""))
    for u in s["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['kim']} {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
