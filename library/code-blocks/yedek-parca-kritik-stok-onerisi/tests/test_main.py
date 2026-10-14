import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def p(kod, kritiklik, tedarik, stok, fiyat=None):
    return {"kod": kod, "ad": "", "kritiklik": kritiklik, "tedarik": D(tedarik), "stok": D(stok), "fiyat": None if fiyat is None else D(fiyat)}


def k(kod, miktar, ekipman="E", t=date(2026, 5, 1)):
    return {"tarih": t, "kod": kod, "miktar": D(miktar), "ekipman": ekipman}


def veri():
    parcalar = [p("P1", "Yüksek", 30, 1, 100), p("P2", "Orta", 60, 5, 10), p("P3", "Düşük", 90, 0), p("P4", "Yüksek", 10, 0, 50), p("P5", "Düşük", 20, 4, 5), p("P6", "Orta", 20, 0)]
    kullanim = [k("P1", 6, "E2"), k("p1", 6, "E1", date(2026, 6, 1))] + [k("P2", 12)] * 3 + [k("P3", 2), k("P9", 3), k("P9", 1)]
    return parcalar, kullanim


class OneriTesti(unittest.TestCase):
    def test_minimum_stok(self):
        r = main.oner(*veri(), gun=360)
        self.assertEqual({x["kod"]: (x["durum"], x["tuketim"], x["kullanim"], x["enb"], main.yuvarla(x["talep"]), x["min"], x["neden"], x["eksik"], x["tutar"]) for x in r["parcalar"]},
                         {"P1": ("Minimumun altında", D(12), 2, D(6), D("1.00"), D(6), "Tek seferde kullanılan en büyük miktar", D(5), D(500)), "P2": ("Minimumun altında", D(36), 3, D(12), D("6.00"), D(12), "Tek seferde kullanılan en büyük miktar", D(7), D(70)),
                          "P3": ("Stok yok", D(2), 1, D(2), D("0.50"), D(1), "Tüketim ve tedarik süresi", D(1), None), "P4": ("Stok yok", D(0), 0, D(0), D("0.00"), D(1), "Kritik parça: en az 1", D(1), D(50)),
                          "P5": ("Hareketsiz", D(0), 0, D(0), D("0.00"), D(0), "Tüketim ve tedarik süresi", D(0), D(0)), "P6": ("Yeterli", D(0), 0, D(0), D("0.00"), D(0), "Tüketim ve tedarik süresi", D(0), None)})
        self.assertEqual([x["kod"] for x in r["parcalar"]], ["P4", "P3", "P1", "P2", "P6", "P5"])
        p1 = next(x for x in r["parcalar"] if x["kod"] == "P1")
        self.assertEqual((p1["son"], p1["ekipman"]), (date(2026, 6, 1), "E1, E2"))
        self.assertEqual(r["ozet"], {"parca": 6, "Stok yok": 2, "Minimumun altında": 2, "Yeterli": 1, "Hareketsiz": 1, "tutar": D(620), "fiyatsiz": 1})
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Bilgi", "Hareketsiz stok", "P5"), ("Orta", "Kullanılan parça listede yok", "P9"), ("Orta", "Minimum stoğun altında", "P2"), ("Orta", "Minimum stoğun altında", "P3"), ("Yüksek", "Kritik parça minimumun altında", "P1"),
                          ("Yüksek", "Kritik parça stokta yok", "P4")])

    def test_emniyet_ve_donem(self):
        m = lambda kritiklik, miktar, tedarik, **kw: main.oner([p("X", kritiklik, tedarik, 99)], [k("X", 1)] * miktar, **kw)["parcalar"][0]["min"]   # noqa: E731
        self.assertEqual((m("Düşük", 100, 36, gun=360), m("Orta", 100, 36, gun=360), m("Yüksek", 100, 36, gun=360), m("Düşük", 100, 36, gun=720), m("Düşük", 100, 36, gun=360, emniyet=(D(0), D(0), D(0))),
                          m("Düşük", 101, 36, gun=360, emniyet=(D(0), D(0), D(0))), m("Orta", 1, 1, gun=360), m("Düşük", 1, 0, gun=360)), (D(13), D(15), D(20), D(7), D(10), D(11), D(1), D(0)))
        self.assertEqual(main.oner([p("X", "Yüksek", 0, 1)], [])["parcalar"][0]["durum"], "Yeterli")
        with self.assertRaises(ValueError):
            main.oner([], [])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "parcalar.csv", ORNEK / "parca_kullanimi.csv", cikti)
            self.assertEqual([(x["kod"], x["durum"], x["tuketim"], x["min"], x["eksik"], x["tutar"]) for x in r["parcalar"]],
                             [("YP-004", "Stok yok", D(0), D(1), D(1), D(140000)), ("YP-003", "Minimumun altında", D(6), D(2), D(1), D(900)), ("YP-001", "Yeterli", D(1), D(1), D(0), D(0)), ("YP-006", "Yeterli", D(1), D(1), D(0), D(0)),
                              ("YP-002", "Yeterli", D(20), D(6), D(0), D(0)), ("YP-005", "Yeterli", D(4), D(1), D(0), D(0)), ("YP-007", "Hareketsiz", D(0), D(0), D(0), D(0))])
            self.assertEqual(r["ozet"], {"parca": 7, "Stok yok": 1, "Minimumun altında": 1, "Yeterli": 4, "Hareketsiz": 1, "tutar": D(140900), "fiyatsiz": 0})
            self.assertEqual([(u["onem"], u["tur"], u["kim"][:6]) for u in r["uyarilar"]],
                             [("Yüksek", "Kritik parça stokta yok", "YP-004"), ("Orta", "Kullanılan parça listede yok", "YP-099"), ("Orta", "Minimum stoğun altında", "YP-003"), ("Bilgi", "Hareketsiz stok", "YP-007")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Stok Önerisi", "Özet", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            py, ky = Path(tmp) / "p.csv", Path(tmp) / "k.csv"
            py.write_text("Stok Kodu;Tanım;Önem;Temin Süresi;Stok;Fiyat\nA1;Pompa;KRİTİK;30;2;1.250,50\nA2;;;15;;\na1;;;5;;\nA3;;Orta;x;;\nA4;;Orta;5;-1;\n;;;5;;\n", encoding="utf-8")
            ky.write_text("Kullanım Tarihi;Malzeme Kodu;Adet;Makine\n01.05.2026;A1;2;M1\nx;A2;1,5;\n;A2;0;\n;A2;y;\n;;1;\n", encoding="utf-8")
            pl, uy = main.oku_parcalar(py)
            self.assertEqual(([(x["kod"], x["ad"], x["kritiklik"], x["tedarik"], x["stok"], x["fiyat"]) for x in pl], [(u["tur"], u["kim"]) for u in uy]),
                             ([("A1", "Pompa", "Yüksek", D(30), D(2), D("1250.50")), ("A2", "", "Orta", D(15), D(0), None)], [("Kritiklik yazılmamış", "A2"), ("Parça okunamadı", "a1"), ("Parça okunamadı", "A3"), ("Parça okunamadı", "A4")]))
            kl, uy = main.oku_kullanim(ky)
            self.assertEqual(([(x["tarih"], x["kod"], x["miktar"], x["ekipman"]) for x in kl], [u["tur"] for u in uy]), ([(date(2026, 5, 1), "A1", D(2), "M1"), (None, "A2", D("1.5"), "")], ["Kullanım kaydı okunamadı"] * 2))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--parcalar", str(py), "--kullanim", str(ky), "--cikti", c]
            self.assertEqual(main.main(arg + ["--gun", "730", "--emniyet", "150;75;0"]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--gun", "0"], ["--emniyet", "100;50"], ["--emniyet", "a;b;c"])] + [main.main(["--parcalar", str(py)])], [2] * 4)
            self.assertEqual(main.main(["--parcalar", str(py), "--kullanim", str(Path(tmp) / "yok.csv")]), 1)
            py.write_text("Parça Kodu;Parça Adı\nA;B\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
