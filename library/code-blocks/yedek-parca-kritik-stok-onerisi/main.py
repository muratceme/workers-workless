"""
Yedek Parça Kritik Stok Önerisi — Workers / Workless kod bloğu
Bakım Onarım › Bakım Mühendisi

Arıza ve bakımda kullanılan parça kayıtları ile tedarik süresinden kritik yedek parçalar için minimum stok seviyesi önerir.
  Günlük tüketim        = dönemde kullanılan miktar ÷ --gun (365)
  Tedarik süresi talebi = günlük tüketim × tedarik süresi (gün)
  Emniyet payı          = tedarik süresi talebi × kritiklik katsayısı   (--emniyet 100;50;25 → Yüksek %100, Orta %50, Düşük %25)
  Minimum stok          = yukarı yuvarla(tedarik süresi talebi + emniyet payı)
      Yüksek ve Orta kritiklikte en az "tek seferde kullanılan en büyük miktar"; Yüksek kritiklikte ayrıca en az 1
  Eksik miktar          = minimum stok − mevcut stok   (sıfırdan büyükse);   sipariş tutarı = eksik × birim fiyat
  Durum                 = Stok yok · Minimumun altında · Yeterli · Hareketsiz (dönemde kullanım yok, stok minimumun üstünde)
Basit bir kural hesabıdır; seyrek kullanılan parçalar için istatistiksel stok modeli değildir. Kritikliği siz belirlersiniz.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
from decimal import ROUND_CEILING

P_SUTUN = {"kod": ("parca kodu", "stok kodu", "malzeme kodu", "kod"), "ad": ("parca adi", "malzeme adi", "tanim", "ad"), "kritiklik": ("kritiklik", "kritiklik derecesi", "onem"),
           "tedarik": ("tedarik suresi gun", "tedarik suresi", "temin suresi"), "stok": ("mevcut stok", "stok", "eldeki stok"), "fiyat": ("birim fiyat", "fiyat", "birim maliyet")}
K_SUTUN = {"tarih": ("tarih", "kullanim tarihi", "is emri tarihi"), "kod": ("parca kodu", "stok kodu", "malzeme kodu", "kod"), "miktar": ("miktar", "kullanilan miktar", "adet"), "ekipman": ("ekipman", "makine")}
KRITIK = {"yuksek": "Yüksek", "kritik": "Yüksek", "a": "Yüksek", "orta": "Orta", "b": "Orta", "dusuk": "Düşük", "c": "Düşük"}
SIRA = ("Yüksek", "Orta", "Düşük")
DURUM = ("Stok yok", "Minimumun altında", "Yeterli", "Hareketsiz")
DURUM_RENK = {"Stok yok": KIRMIZI, "Minimumun altında": SARI, "Yeterli": YESIL, "Hareketsiz": MAVI}
SIFIR = Decimal(0)


def oku_parcalar(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, P_SUTUN, ("kod", "tedarik")):
        kod = metin(r["kod"])
        if not kod:
            continue
        tedarik, stok, fiyat = para(r["tedarik"]), para(r["stok"]) if metin(r["stok"]) else SIFIR, para(r["fiyat"]) if metin(r["fiyat"]) else None
        if tedarik is None or tedarik < 0 or stok is None or stok < 0 or (metin(r["fiyat"]) and (fiyat is None or fiyat < 0)) or katla(kod) in gorulen:
            uy.append(uyari("Yüksek", "Parça okunamadı", kod, f"{yol.name} satır {r['_satir']}: " + ("aynı parça kodu daha önce geçti" if katla(kod) in gorulen else "tedarik süresi, stok ya da fiyat okunamadı") + ". Satır alınmadı."))
            continue
        gorulen.add(katla(kod))
        kritiklik = KRITIK.get(katla(r["kritiklik"]))
        if kritiklik is None:
            uy.append(uyari("Orta", "Kritiklik yazılmamış", kod, f"Kritiklik '{metin(r['kritiklik'])}' Yüksek, Orta ya da Düşük olarak okunamadı; Orta sayıldı."))
        l.append({"kod": kod, "ad": metin(r["ad"]), "kritiklik": kritiklik or "Orta", "tedarik": tedarik, "stok": stok, "fiyat": fiyat})
    return l, uy


def oku_kullanim(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, K_SUTUN, ("kod", "miktar")):
        kod = metin(r["kod"])
        if not kod:
            continue
        miktar = para(r["miktar"])
        if miktar is None or miktar <= 0:
            uy.append(uyari("Yüksek", "Kullanım kaydı okunamadı", kod, f"{yol.name} satır {r['_satir']}: miktar '{metin(r['miktar'])}' okunamadı. Satır alınmadı."))
            continue
        l.append({"tarih": tarih(r["tarih"]), "kod": kod, "miktar": miktar, "ekipman": metin(r["ekipman"])})
    return l, uy


def oner(parcalar: list[dict], kullanim: list[dict], gun: int = 365, emniyet: tuple[Decimal, ...] = (Decimal(100), Decimal(50), Decimal(25))) -> dict:
    if not parcalar:
        raise ValueError("Parça listesi boş.")
    kat, uy, grup = dict(zip(SIRA, emniyet)), [], {}
    for k in kullanim:
        grup.setdefault(katla(k["kod"]), []).append(k)
    for p in parcalar:
        kl = grup.pop(katla(p["kod"]), [])
        p["kullanim"], p["tuketim"], p["enb"] = len(kl), sum((k["miktar"] for k in kl), SIFIR), max((k["miktar"] for k in kl), default=SIFIR)
        p["son"] = max((k["tarih"] for k in kl if k["tarih"]), default=None)
        p["ekipman"] = ", ".join(sorted({k["ekipman"] for k in kl if k["ekipman"]}, key=katla))
        p["gunluk"] = p["tuketim"] / gun
        p["talep"] = p["gunluk"] * p["tedarik"]
        hesap = (p["talep"] * (1 + kat[p["kritiklik"]] / 100)).to_integral_value(ROUND_CEILING)
        taban = SIFIR if p["kritiklik"] == "Düşük" else max(p["enb"].to_integral_value(ROUND_CEILING), Decimal(1) if p["kritiklik"] == "Yüksek" else SIFIR)
        p["min"], p["neden"] = max(hesap, taban), "Tüketim ve tedarik süresi" if hesap >= taban else "Tek seferde kullanılan en büyük miktar" if p["enb"] >= 1 else "Kritik parça: en az 1"
        p["eksik"] = max(p["min"] - p["stok"], SIFIR)
        p["tutar"] = p["eksik"] * p["fiyat"] if p["fiyat"] is not None else None
        p["durum"] = "Stok yok" if p["stok"] == 0 and p["min"] > 0 else "Minimumun altında" if p["stok"] < p["min"] else "Hareketsiz" if p["tuketim"] == 0 and p["stok"] > p["min"] else "Yeterli"
        kim = f"{p['kod']} · {p['ad']}" if p["ad"] else p["kod"]
        bilgi = f"Önerilen minimum {p['min']:f}, mevcut {p['stok'].normalize():f}, tedarik süresi {p['tedarik'].normalize():f} gün" + (f"; sipariş tutarı {p['tutar']:,.2f}" if p["tutar"] else "") + "."
        if p["durum"] in ("Stok yok", "Minimumun altında") and p["kritiklik"] == "Yüksek":
            uy.append(uyari("Yüksek", "Kritik parça minimumun altında" if p["stok"] > 0 else "Kritik parça stokta yok", kim, bilgi))
        elif p["durum"] in ("Stok yok", "Minimumun altında"):
            uy.append(uyari("Orta", "Minimum stoğun altında", kim, bilgi))
        elif p["durum"] == "Hareketsiz":
            uy.append(uyari("Bilgi", "Hareketsiz stok", kim, f"Dönemde hiç kullanılmadı; stok {p['stok'].normalize():f}" + (f", değeri {p['stok'] * p['fiyat']:,.2f}" if p["fiyat"] else "") + ". Parçanın hâlâ gerekli olup olmadığını değerlendirin."))
    for kl in grup.values():
        uy.append(uyari("Orta", "Kullanılan parça listede yok", kl[0]["kod"], f"{len(kl)} kullanım kaydı var ({sum(k['miktar'] for k in kl).normalize():f} adet) ama parça listesinde tanımlı değil; öneri üretilemedi."))
    parcalar.sort(key=lambda p: (DURUM.index(p["durum"]), SIRA.index(p["kritiklik"]), -(p["tutar"] or 0), katla(p["kod"])))
    return {"parcalar": parcalar, "ozet": {"parca": len(parcalar), **{d: sum(1 for p in parcalar if p["durum"] == d) for d in DURUM}, "tutar": sum((p["tutar"] or SIFIR for p in parcalar), SIFIR),
                                           "fiyatsiz": sum(1 for p in parcalar if p["eksik"] > 0 and p["fiyat"] is None)}, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, gun: int, emniyet: tuple[Decimal, ...]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Stok Önerisi"
    _baslik(ws, ["Durum", "Parça Kodu", "Parça Adı", "Kritiklik", "Dönem Tüketimi", "Kullanım Sayısı", "Tek Seferde En Çok", "Tedarik Süresi (gün)", "Tedarik Süresi Talebi", "Önerilen Minimum Stok", "Minimumu Belirleyen", "Mevcut Stok",
                 "Eksik Miktar", "Birim Fiyat", "Sipariş Tutarı", "Son Kullanım", "Kullanıldığı Ekipman"], (18, 13, 28, 10, 10, 10, 10, 10, 11, 11, 34, 10, 9, 13, 15, 12, 30))
    for p in r["parcalar"]:
        _satir(ws, [p["durum"], p["kod"], p["ad"], p["kritiklik"], p["tuketim"] or None, p["kullanim"] or None, p["enb"] or None, p["tedarik"], yuvarla(p["talep"]), p["min"], p["neden"], p["stok"], p["eksik"] or None, p["fiyat"], p["tutar"] or None,
                    p["son"], p["ekipman"]], para_s=(9, 14, 15), tarih_s=(16,), dolgu={1: DURUM_RENK[p["durum"]], 4: KIRMIZI if p["kritiklik"] == "Yüksek" else None})
    ws = wb.create_sheet("Özet")
    o = r["ozet"]
    _baslik(ws, ["Gösterge", "Değer"], (44, 16))
    for etiket, deger, para_mi in [("Parça", o["parca"], False)] + [(d, o[d], False) for d in DURUM] + [("Minimuma tamamlamak için sipariş tutarı", o["tutar"], True), ("Eksiği olup fiyatı yazılmamış parça", o["fiyatsiz"], False)]:
        _satir(ws, [etiket, deger], para_s=(2,) if para_mi else ())
    ws.append([])
    ws.append([f"Dönem {gun} gün · emniyet payı: Yüksek %{emniyet[0]}, Orta %{emniyet[1]}, Düşük %{emniyet[2]}"])
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(p_yol: Path, k_yol: Path, cikti: Path, gun: int = 365, emniyet: tuple[Decimal, ...] = (Decimal(100), Decimal(50), Decimal(25))) -> dict:
    parcalar, uy1 = oku_parcalar(p_yol)
    kullanim, uy2 = oku_kullanim(k_yol)
    r = oner(parcalar, kullanim, gun, emniyet)
    r["uyarilar"] = sirala(uy1 + uy2 + r["uyarilar"])
    rapor_yaz(cikti, r, gun, emniyet)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Arıza geçmişi ve tedarik süresine göre kritik yedek parçaların minimum stok seviyesini önerir.")
    p.add_argument("--parcalar", type=Path, help="Parça Kodu, Parça Adı, Kritiklik, Tedarik Süresi (gün), Mevcut Stok, Birim Fiyat")
    p.add_argument("--kullanim", type=Path, help="Tarih, Parça Kodu, Miktar, Ekipman (arıza ve bakımda kullanılan parçalar)")
    p.add_argument("--gun", type=int, default=365, help="Kullanım kayıtlarının kapsadığı dönem, gün (varsayılan 365)")
    p.add_argument("--emniyet", default="100;50;25", help="Yüksek;Orta;Düşük kritiklik için emniyet payı, tedarik süresi talebinin yüzdesi (varsayılan 100;50;25)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "yedek_parca_kritik_stok.xlsx")
    a = p.parse_args(argv)
    if (a.parcalar is None) != (a.kullanim is None):
        print("[X] --parcalar ve --kullanim birlikte verilmeli")
        return 2
    p_yol, k_yol = a.parcalar or ORNEK / "parcalar.csv", a.kullanim or ORNEK / "parca_kullanimi.csv"
    emniyet = tuple(para(x) for x in a.emniyet.split(";"))
    if a.gun < 1 or len(emniyet) != 3 or any(x is None or x < 0 for x in emniyet):
        print("[X] --gun en az 1; --emniyet üç yüzde olmalı (ör. 100;50;25)")
        return 2
    for yol in (p_yol, k_yol):
        if not yol.exists():
            print(f"[X] Dosya bulunamadı: {yol}")
            return 1
    try:
        r = calistir(p_yol, k_yol, a.cikti, a.gun, emniyet)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {o['parca']} parça · stok yok {o['Stok yok']} · minimumun altında {o['Minimumun altında']} · hareketsiz {o['Hareketsiz']} · sipariş tutarı {o['tutar']:,.2f}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
