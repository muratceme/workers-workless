"""
Log Anomali Tespiti — Workers / Workless kod bloğu
Bilgi Teknolojileri › Sistem Uzmanı

Sunucu ve uygulama loglarında olağandışı örüntüleri bulur ve özetler. Biçim satırlardan tanınır:
  - Web erişim logu (Apache/Nginx "combined"/"common"): IP bazında yoğun 401/403 (parola deneme) ve ardından
    başarılı giriş, çok sayıda farklı 404 (dizin/zafiyet taraması), saldırı imzaları (SQL enjeksiyonu, XSS,
    dizin aşma, .env/.git gibi hassas dosyalar, Log4Shell), dakikada aşırı istek, 5xx artışı, mesai dışı
    yönetim paneli erişimi.
  - Linux auth.log / secure (sshd, sudo): IP bazında başarısız giriş yoğunluğu, var olmayan kullanıcı denemeleri,
    deneme sonrası başarılı giriş (kritik), root ile parola girişi, kullanıcının yeni IP'den girişi, sudo parola
    hataları, mesai dışı giriş.
  - Uygulama logu ("2026-10-07 14:30:01,123 ERROR [modül] mesaj"): hata sayısındaki ani artış ve kısa sürede ortaya
    çıkan yeni hata türleri (mesajdaki sayı, kimlik ve adresler normalleştirilerek imza çıkarılır).
Ani artış: zaman dilimi sayısı, tüm dilimlerin medyanı + 4 × (1,4826 × MAD) eşiğini ve en az --min-artis değerini
aşarsa işaretlenir. İnternete bağlanmaz; logları hiçbir yere göndermez.

Kullanım:
    python main.py                                                  # örnek loglar (07.10.2026)
    python main.py --girdi /var/log/nginx/access.log /var/log/auth.log uygulama.log --yil 2026
    python main.py --girdi ./loglar --pencere 10 --esik-giris 8 --mesai 08-19
"""
from __future__ import annotations

import argparse
import gzip
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from statistics import median

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
AYLAR = {a: i for i, a in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}
AYLAR.update({"oca": 1, "sub": 2, "nis": 4, "haz": 6, "tem": 7, "agu": 8, "eyl": 9, "eki": 10, "kas": 11, "ara": 12})

ERISIM = re.compile(r'^(?P<ip>[\da-fA-F:.]+) \S+ (?P<kul>\S+) \[(?P<t>[^\]]+)\] "(?P<m>[A-Z]+) (?P<yol>\S+)[^"]*" (?P<kod>\d{3}) (?P<boy>\S+)'
                    r'(?: "(?P<ref>[^"]*)" "(?P<ua>[^"]*)")?')
SYSLOG = re.compile(r"^(?:(?P<ay>[A-Za-z]{3})\s+(?P<gun>\d{1,2}) (?P<saat>\d\d:\d\d:\d\d)|(?P<iso>\d{4}-\d\d-\d\dT\d\d:\d\d:\d\d)\S*) "
                    r"(?P<host>\S+) (?P<prog>[\w\-/.()]+?)(?:\[\d+\])?: (?P<msg>.*)$")
UYGULAMA = re.compile(r"^(?P<t>\d{4}-\d\d-\d\d[ T]\d\d:\d\d:\d\d)(?:[.,]\d+)?(?:Z|[+-]\d\d:?\d\d)?\s+\[?(?P<sev>TRACE|DEBUG|INFO|NOTICE|WARN|WARNING|"
                      r"ERROR|ERR|FATAL|CRITICAL|SEVERE)\]?\s+(?P<msg>.*)$", re.I)
HATA_SEVIYE = {"ERROR", "ERR", "FATAL", "CRITICAL", "SEVERE"}

BASARISIZ = re.compile(r"Failed (?:password|publickey|keyboard-interactive/pam) for (?P<inv>invalid user )?(?P<kul>\S+) from (?P<ip>\S+)")
BASARILI = re.compile(r"Accepted (?P<yontem>\S+) for (?P<kul>\S+) from (?P<ip>\S+)")
SUDO_HATA = re.compile(r"(?:sudo|su).*authentication failure.*\buser=(?P<kul>\S+)")

IMZALAR = [
    ("SQL enjeksiyonu", re.compile(r"union\s+(all\s+)?select|'\s*or\s+'?\d+'?\s*=\s*'?\d|\bsleep\s*\(\d|information_schema|;\s*drop\s+table", re.I)),
    ("XSS", re.compile(r"<script|javascript:|onerror\s*=|onload\s*=", re.I)),
    ("Dizin aşma", re.compile(r"\.\./|/etc/passwd|\\windows\\win\.ini", re.I)),
    ("Log4Shell", re.compile(r"\$\{jndi:", re.I)),
    ("Komut çalıştırma", re.compile(r"cmd\.exe|/bin/(ba)?sh|;\s*(wget|curl)\s", re.I)),
    ("Hassas dosya / panel arama", re.compile(r"/\.env\b|/\.git/|wp-config|/phpmyadmin|/xmlrpc\.php|/wp-login\.php|/server-status|"
                                              r"\.(bak|sql|zip)$", re.I)),
]
ISLENMIS_TEHLIKELI = {"SQL enjeksiyonu", "XSS", "Dizin aşma", "Log4Shell", "Komut çalıştırma"}


def unquote_plus(s: str) -> str:
    """URL yüzde kodlamasını çözer (%27 → ', + → boşluk); ağ modülü kullanılmaz."""
    b = re.sub(rb"%([0-9A-Fa-f]{2})", lambda m: bytes([int(m.group(1), 16)]), s.replace("+", " ").encode("utf-8", "surrogateescape"))
    return b.decode("utf-8", "replace")


@dataclass
class Olay:
    dosya: str
    tur: str                  # erisim | auth | uygulama
    zaman: datetime
    ip: str = ""
    kullanici: str = ""
    yontem: str = ""
    yol: str = ""
    kod: int = 0
    seviye: str = ""
    mesaj: str = ""
    satir: str = ""


@dataclass
class Bulgu:
    onem: str
    tur: str
    kaynak: str
    varlik: str
    bas: datetime | None
    bit: datetime | None
    adet: int
    aciklama: str
    ornek: str = ""


# ----------------------------------------------------------------------------
# Okuma
# ----------------------------------------------------------------------------

def satirlari_oku(yol: Path) -> list[str]:
    ham = gzip.open(yol, "rb").read() if yol.suffix == ".gz" else yol.read_bytes()
    for kod in ("utf-8", "cp1254", "latin-1"):
        try:
            return ham.decode(kod).splitlines()
        except UnicodeDecodeError:
            continue
    return []


def erisim_zamani(t: str) -> datetime | None:
    m = re.match(r"(\d{1,2})/(\w{3})/(\d{4}):(\d\d):(\d\d):(\d\d)", t)
    if not m or m.group(2).lower()[:3] not in AYLAR:
        return None
    return datetime(int(m.group(3)), AYLAR[m.group(2).lower()[:3]], int(m.group(1)), int(m.group(4)), int(m.group(5)), int(m.group(6)))


def satir_coz(satir: str, dosya: str, yil: int) -> Olay | None:
    m = ERISIM.match(satir)
    if m:
        z = erisim_zamani(m.group("t"))
        if z:
            return Olay(dosya, "erisim", z, ip=m.group("ip"), kullanici="" if m.group("kul") == "-" else m.group("kul"), yontem=m.group("m"),
                        yol=m.group("yol"), kod=int(m.group("kod")), satir=satir)
    m = UYGULAMA.match(satir)
    if m:
        return Olay(dosya, "uygulama", datetime.fromisoformat(m.group("t").replace("T", " ")), seviye=m.group("sev").upper(),
                    mesaj=m.group("msg"), satir=satir)
    m = SYSLOG.match(satir)
    if m:
        if m.group("iso"):
            z = datetime.fromisoformat(m.group("iso"))
        else:
            ay = AYLAR.get(m.group("ay").lower())
            if not ay:
                return None
            z = datetime.strptime(f"{yil}-{ay:02d}-{int(m.group('gun')):02d} {m.group('saat')}", "%Y-%m-%d %H:%M:%S")
        return Olay(dosya, "auth", z, mesaj=m.group("msg"), yontem=m.group("prog"), satir=satir)
    return None


def loglari_oku(yollar: list[Path], yil: int) -> tuple[list[Olay], dict[str, dict]]:
    dosyalar = []
    for y in yollar:
        dosyalar += sorted(p for p in y.rglob("*") if p.is_file() and not p.name.startswith(".")) if y.is_dir() else [y]
    olaylar, ozet = [], {}
    for y in dosyalar:
        satirlar = satirlari_oku(y)
        cozulen = [o for o in (satir_coz(s, y.name, yil) for s in satirlar if s.strip()) if o]
        tur = Counter(o.tur for o in cozulen).most_common(1)
        ozet[y.name] = {"satir": sum(1 for s in satirlar if s.strip()), "cozulen": len(cozulen), "tur": tur[0][0] if tur else "-",
                        "ilk": min((o.zaman for o in cozulen), default=None), "son": max((o.zaman for o in cozulen), default=None)}
        olaylar += cozulen
    olaylar.sort(key=lambda o: o.zaman)
    return olaylar, ozet


# ----------------------------------------------------------------------------
# Tespit
# ----------------------------------------------------------------------------

def en_yogun(zamanlar: list[datetime], pencere: timedelta) -> int:
    """Kayan pencerede en fazla olay sayısı."""
    en, j = 0, 0
    for i, z in enumerate(zamanlar):
        while z - zamanlar[j] > pencere:
            j += 1
        en = max(en, i - j + 1)
    return en


def mesai_disi(z: datetime, mesai: tuple[int, int]) -> bool:
    return z.weekday() >= 5 or not (mesai[0] <= z.hour < mesai[1])


def giris_denemeleri(olaylar: list[Olay], p: dict) -> list[Bulgu]:
    bulgular = []
    hatali, basarili = defaultdict(list), []
    for o in olaylar:
        if o.tur != "auth":
            continue
        m = BASARISIZ.search(o.mesaj)          # "Invalid user" satırı aynı denemenin ön kaydıdır; çift sayılmaz
        if m:
            hatali[m.group("ip")].append((o, m.group("kul"), bool(m.group("inv"))))
            continue
        m = BASARILI.search(o.mesaj)
        if m:
            basarili.append((o, m.group("yontem"), m.group("kul"), m.group("ip")))
    for ip, ks in hatali.items():
        zaman = [k[0].zaman for k in ks]
        yogun = en_yogun(zaman, p["pencere"])
        sonra = [b for b in basarili if b[3] == ip and b[0].zaman >= zaman[0]]
        if yogun >= p["esik_giris"] or sonra:
            kullanicilar = Counter(k[1] for k in ks)
            gecersiz = sum(1 for k in ks if k[2])
            ac = (f"{len(ks)} başarısız SSH girişi ({p['pencere_dk']} dk içinde en fazla {yogun}); denenen kullanıcılar: "
                  + ", ".join(f"{u} ({n})" for u, n in kullanicilar.most_common(6)))
            if gecersiz >= 5:
                ac += f". {gecersiz} deneme var olmayan kullanıcı adlarıyla (kullanıcı adı tahmini)"
            onem = "Yüksek" if yogun >= p["esik_giris"] else "Dikkat"
            if sonra:
                b = sonra[0]
                onem, ac = "Kritik", ac + f". ARDINDAN {b[0].zaman:%d.%m %H:%M:%S} başarılı giriş: {b[2]} ({b[1]}) — hesabı ve oturumu hemen inceleyin"
            bulgular.append(Bulgu(onem, "SSH parola deneme", ks[0][0].dosya, ip, zaman[0], (sonra[0][0].zaman if sonra else zaman[-1]),
                                  len(ks), ac, ks[0][0].satir))
    gorulen_ip = defaultdict(list)
    for o, yontem, kul, ip in basarili:
        if kul == "root" and yontem == "password":
            bulgular.append(Bulgu("Yüksek", "root ile parola girişi", o.dosya, ip, o.zaman, o.zaman, 1,
                                  "root hesabına parolayla uzaktan giriş yapıldı (PermitRootLogin ve anahtar kullanımı gözden geçirilmeli)", o.satir))
        if gorulen_ip[kul] and ip not in gorulen_ip[kul]:
            bulgular.append(Bulgu("Bilgi", "Yeni IP'den giriş", o.dosya, f"{kul} @ {ip}", o.zaman, o.zaman, 1,
                                  f"{kul} bu logda daha önce {', '.join(sorted(set(gorulen_ip[kul])))} adresinden girmişti", o.satir))
        gorulen_ip[kul].append(ip)
        if mesai_disi(o.zaman, p["mesai"]) and not (kul == "root" and yontem == "password"):
            bulgular.append(Bulgu("Bilgi", "Mesai dışı giriş", o.dosya, f"{kul} @ {ip}", o.zaman, o.zaman, 1, f"{o.zaman:%H:%M} başarılı SSH girişi ({yontem})",
                                  o.satir))
    sudo = defaultdict(list)
    for o in olaylar:
        if o.tur == "auth" and (m := SUDO_HATA.search(o.mesaj)):
            sudo[m.group("kul")].append(o)
    for kul, os_ in sudo.items():
        if len(os_) >= 3:
            bulgular.append(Bulgu("Dikkat", "sudo parola hatası", os_[0].dosya, kul, os_[0].zaman, os_[-1].zaman, len(os_),
                                  f"{kul} kullanıcısının {len(os_)} sudo/su kimlik doğrulama hatası", os_[0].satir))
    return bulgular


def web_analizi(olaylar: list[Olay], p: dict) -> list[Bulgu]:
    bulgular = []
    erisim = [o for o in olaylar if o.tur == "erisim"]
    ip_ol = defaultdict(list)
    for o in erisim:
        ip_ol[o.ip].append(o)
    for ip, os_ in ip_ol.items():
        red = [o for o in os_ if o.kod in (401, 403)]
        if red:
            yogun = en_yogun([o.zaman for o in red], p["pencere"])
            if yogun >= p["esik_giris"]:
                yollar = Counter(o.yol.split("?")[0] for o in red)
                hedef = {y for y, _ in yollar.most_common(3)}
                basari = next((o for o in os_ if o.zaman > red[0].zaman and 200 <= o.kod < 400 and o.yol.split("?")[0] in hedef), None)
                ac = f"{len(red)} adet 401/403 ({p['pencere_dk']} dk içinde en fazla {yogun}); yollar: " + ", ".join(f"{y} ({n})" for y, n in yollar.most_common(3))
                onem = "Yüksek"
                if basari:
                    onem, ac = "Kritik", ac + f". ARDINDAN {basari.zaman:%H:%M:%S} {basari.yontem} {basari.yol} → {basari.kod}: deneme başarılı olmuş olabilir"
                bulgular.append(Bulgu(onem, "Web parola deneme", red[0].dosya, ip, red[0].zaman, (basari or red[-1]).zaman, len(red), ac, red[0].satir))
        bulunamadi = {o.yol for o in os_ if o.kod == 404}
        if len(bulunamadi) >= p["esik_404"]:
            n404 = [o for o in os_ if o.kod == 404]
            bulgular.append(Bulgu("Yüksek", "Dizin / zafiyet taraması", n404[0].dosya, ip, n404[0].zaman, n404[-1].zaman, len(n404),
                                  f"{len(bulunamadi)} farklı yol 404 döndü; örnek: " + ", ".join(sorted(bulunamadi)[:5]), n404[0].satir))
        dakika = Counter(o.zaman.replace(second=0) for o in os_)
        tepe = max(dakika.values())
        if tepe >= p["esik_istek"]:
            d = max(dakika, key=dakika.get)
            bulgular.append(Bulgu("Dikkat", "Aşırı istek", os_[0].dosya, ip, d, d + timedelta(minutes=1), tepe,
                                  f"{d:%H:%M} dakikasında {tepe} istek (bot, kazıma veya DoS olabilir)"))
    imza = defaultdict(list)
    for o in erisim:
        yol = unquote_plus(o.yol)
        for ad, desen in IMZALAR:
            if desen.search(yol):
                imza[(o.ip, ad)].append(o)
                break
    for (ip, ad), os_ in imza.items():
        basarili = [o for o in os_ if 200 <= o.kod < 300]
        onem = "Yüksek" if ad in ISLENMIS_TEHLIKELI and (basarili or any(o.kod >= 500 for o in os_)) else "Dikkat"
        ac = f"{len(os_)} istek; yanıt kodları: " + ", ".join(f"{k} ({n})" for k, n in Counter(o.kod for o in os_).most_common())
        if onem == "Yüksek":
            ac += ". 2xx/5xx yanıt: istek uygulamaya ulaşmış; uygulama logunu ve girdi doğrulamasını kontrol edin"
        bulgular.append(Bulgu(onem, f"Saldırı imzası: {ad}", os_[0].dosya, ip, os_[0].zaman, os_[-1].zaman, len(os_), ac,
                              unquote_plus(os_[0].yol)[:200]))
    yonetim = re.compile(p["yonetim_yolu"], re.I)
    panel = defaultdict(list)
    for o in erisim:
        if 200 <= o.kod < 300 and yonetim.search(o.yol) and mesai_disi(o.zaman, p["mesai"]):
            panel[o.ip].append(o)
    for ip, os_ in panel.items():
        bulgular.append(Bulgu("Dikkat", "Mesai dışı yönetim erişimi", os_[0].dosya, ip, os_[0].zaman, os_[-1].zaman, len(os_),
                              "Başarılı yönetim paneli istekleri: " + ", ".join(sorted({f"{o.yontem} {o.yol}" for o in os_})[:4]), os_[0].satir))
    return bulgular


def dilimler(olaylar: list[Olay], pencere: timedelta) -> list[datetime]:
    if not olaylar:
        return []
    dk = int(pencere.total_seconds() // 60)
    bas = min(o.zaman for o in olaylar).replace(second=0, microsecond=0)
    bas -= timedelta(minutes=bas.minute % dk)
    son = max(o.zaman for o in olaylar)
    sonuc = []
    while bas <= son:
        sonuc.append(bas)
        bas += pencere
    return sonuc


def dilim_anahtari(z: datetime, pencere: timedelta) -> datetime:
    dk = int(pencere.total_seconds() // 60)
    z = z.replace(second=0, microsecond=0)
    return z - timedelta(minutes=z.minute % dk)


def ani_artis(sayac: Counter, dl: list[datetime], pencere: timedelta, min_artis: int) -> tuple[list[tuple[datetime, datetime, int, float]], float]:
    """Medyan + 4 × 1,4826 × MAD eşiğini aşan ardışık dilimler: (başlangıç, bitiş, toplam, eşik)."""
    degerler = [sayac.get(d, 0) for d in dl]
    if len(degerler) < 6:
        return [], 0.0
    med = median(degerler)
    mad = median(abs(x - med) for x in degerler)
    esik = max(min_artis, med + 4 * 1.4826 * mad if mad else max(med * 3, 1))
    sonuc, aktif = [], None
    for d, x in zip(dl, degerler):
        if x >= esik:
            aktif = [aktif[0], d + pencere, aktif[2] + x] if aktif else [d, d + pencere, x]
        elif aktif:
            sonuc.append((*aktif, esik))
            aktif = None
    if aktif:
        sonuc.append((*aktif, esik))
    return sonuc, med


def imza(mesaj: str) -> str:
    s = re.sub(r"https?://([^/\s]+)\S*", r"<url:\1>", mesaj)
    s = re.sub(r"\b[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}\b", "<id>", s, flags=re.I)
    s = re.sub(r"\b0x[0-9a-f]+\b|\b[0-9a-f]{16,}\b", "<hex>", s, flags=re.I)
    s = re.sub(r"'[^']*'|\"[^\"]*\"", "<s>", s)
    s = re.sub(r"\d+(?:[.,]\d+)*", "<n>", s)
    return " ".join(s.split())[:160]


def zaman_analizi(olaylar: list[Olay], p: dict) -> tuple[list[Bulgu], dict, dict]:
    bulgular = []
    pn = p["pencere"]
    seriler = {"İstek": Counter(), "4xx": Counter(), "5xx": Counter(), "Başarısız giriş": Counter(), "Uygulama hatası": Counter()}
    for o in olaylar:
        d = dilim_anahtari(o.zaman, pn)
        if o.tur == "erisim":
            seriler["İstek"][d] += 1
            if 400 <= o.kod < 500:
                seriler["4xx"][d] += 1
            elif o.kod >= 500:
                seriler["5xx"][d] += 1
        elif o.tur == "auth" and BASARISIZ.search(o.mesaj):
            seriler["Başarısız giriş"][d] += 1
        elif o.tur == "uygulama" and o.seviye in HATA_SEVIYE:
            seriler["Uygulama hatası"][d] += 1
    dl = dilimler(olaylar, pn)
    for ad, kaynak in (("5xx", "erisim"), ("Uygulama hatası", "uygulama")):
        tur_ol = [o for o in olaylar if o.tur == kaynak]
        if not tur_ol:
            continue
        aktif = dilimler(tur_ol, pn)
        for bas, bit, toplam, esik in ani_artis(seriler[ad], aktif, pn, p["min_artis"])[0]:
            ornekler = [o for o in tur_ol if bas <= o.zaman < bit and ((o.kod >= 500) if kaynak == "erisim" else o.seviye in HATA_SEVIYE)]
            detay = Counter(str(o.kod) if kaynak == "erisim" else imza(o.mesaj) for o in ornekler).most_common(3)
            bulgular.append(Bulgu("Yüksek", f"{ad} artışı", ornekler[0].dosya if ornekler else "", "-", bas, bit, toplam,
                                  f"{bas:%H:%M}-{bit:%H:%M} arasında {toplam} ({p['pencere_dk']} dk'lık dilim eşiği {esik:.0f}); en sık: "
                                  + "; ".join(f"{k} ({n})" for k, n in detay), ornekler[0].satir if ornekler else ""))
    # Hata imzaları
    hatalar = [o for o in olaylar if o.tur == "uygulama" and o.seviye in HATA_SEVIYE]
    imzalar = defaultdict(list)
    for o in hatalar:
        imzalar[imza(o.mesaj)].append(o)
    if hatalar:
        sure = max(o.zaman for o in hatalar) - min(o.zaman for o in hatalar)
        for s, os_ in imzalar.items():
            yayilim = os_[-1].zaman - os_[0].zaman
            if len(os_) >= 5 and sure >= timedelta(hours=6) and yayilim <= timedelta(hours=1):
                bulgular.append(Bulgu("Dikkat", "Ani ortaya çıkan hata türü", os_[0].dosya, s[:60], os_[0].zaman, os_[-1].zaman, len(os_),
                                      f"{len(os_)} kez, yalnız {os_[0].zaman:%H:%M}-{os_[-1].zaman:%H:%M} arasında görüldü (log süresi "
                                      f"{sure.total_seconds() / 3600:.0f} saat)", os_[0].satir))
    return bulgular, {"dilimler": dl, "seriler": seriler}, imzalar


def analiz_et(olaylar: list[Olay], p: dict) -> tuple[list[Bulgu], dict, dict]:
    b = giris_denemeleri(olaylar, p) + web_analizi(olaylar, p)
    zb, zaman, imzalar = zaman_analizi(olaylar, p)
    b += zb
    sira = {"Kritik": 0, "Yüksek": 1, "Dikkat": 2, "Bilgi": 3}
    b.sort(key=lambda x: (sira[x.onem], x.bas or datetime.min))
    return b, zaman, imzalar


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Kritik": "F8C9C6", "Yüksek": "FDE2E1", "Dikkat": "FFF4CE", "Bilgi": "E8F0FE"}
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def rapor_yaz(cikti: Path, olaylar: list[Olay], dosya_ozet: dict, bulgular: list[Bulgu], zaman: dict, imzalar: dict, p: dict) -> None:
    wb = Workbook()
    o = wb.active
    o.title = "Özet"
    _baslik(o, ["Dosya", "Tür", "Satır", "Çözülen", "İlk", "Son"], (28, 12, 10, 10, 18, 18))
    tur_ad = {"erisim": "Web erişim", "auth": "auth / syslog", "uygulama": "Uygulama", "-": "tanınmadı"}
    for ad, x in dosya_ozet.items():
        o.append([ad, tur_ad[x["tur"]], x["satir"], x["cozulen"], x["ilk"], x["son"]])
        for c in (5, 6):
            o.cell(o.max_row, c).number_format = "DD.MM.YYYY HH:MM"
        if x["cozulen"] < x["satir"]:
            o.cell(o.max_row, 4).fill = PatternFill("solid", fgColor=RENK["Dikkat"])
    o.append([])
    _baslik(o, ["Önem", "Bulgu sayısı"], ())
    say = Counter(b.onem for b in bulgular)
    for on in RENK:
        o.append([on, say[on]])
        o.cell(o.max_row, 1).fill = PatternFill("solid", fgColor=RENK[on])
    o.append([])
    o.append(["Ayarlar", f"pencere {p['pencere_dk']} dk · giriş eşiği {p['esik_giris']} · 404 eşiği {p['esik_404']} · dakikada istek eşiği "
                         f"{p['esik_istek']} · mesai {p['mesai'][0]:02d}-{p['mesai'][1]:02d}"])
    o.append(["Not", "Bulgular olası anomalilerdir; olay müdahalesi için ilgili sistemlerde doğrulanmalıdır. Loglar kişisel veri (IP, kullanıcı "
                     "adı) içerebilir; raporu yetkisiz kişilerle paylaşmayın."])
    o.cell(o.max_row, 2).alignment = UST

    b = wb.create_sheet("Bulgular")
    _baslik(b, ["Önem", "Tür", "Kaynak", "IP / Kullanıcı", "Başlangıç", "Bitiş", "Adet", "Açıklama", "Örnek Satır", "İnceleme Notu"],
            (8, 26, 12, 22, 16, 16, 7, 70, 60, 24))
    for x in bulgular:
        b.append([x.onem, x.tur, x.kaynak, x.varlik, x.bas, x.bit, x.adet, x.aciklama, x.ornek, ""])
        b.cell(b.max_row, 1).fill = PatternFill("solid", fgColor=RENK[x.onem])
        for c in (5, 6):
            b.cell(b.max_row, c).number_format = "DD.MM HH:MM:SS"
        b.cell(b.max_row, 10).fill = PatternFill("solid", fgColor="FFF4CE")
        for h in b[b.max_row]:
            h.alignment = UST
    b.freeze_panes = "B2"
    b.auto_filter.ref = b.dimensions

    z = wb.create_sheet("Zaman Çizelgesi")
    adlar = [a for a, s in zaman["seriler"].items() if s]
    _baslik(z, ["Dilim"] + adlar, (16,) + (14,) * len(adlar))
    for d in zaman["dilimler"]:
        z.append([d] + [zaman["seriler"][a].get(d, 0) for a in adlar])
        z.cell(z.max_row, 1).number_format = "DD.MM HH:MM"
    if zaman["dilimler"] and adlar:
        g = LineChart()
        g.title, g.height, g.width = f"{p['pencere_dk']} dakikalık dilimler", 9, 26
        g.add_data(Reference(z, min_col=2, max_col=1 + len(adlar), min_row=1, max_row=z.max_row), titles_from_data=True)
        g.set_categories(Reference(z, min_col=1, min_row=2, max_row=z.max_row))
        z.add_chart(g, f"{get_column_letter(len(adlar) + 3)}2")

    ip = wb.create_sheet("IP Özeti")
    _baslik(ip, ["IP", "Web İstek", "4xx", "5xx", "Farklı 404 Yolu", "Başarısız SSH", "Başarılı SSH", "İlk", "Son"], (20, 10, 8, 8, 14, 13, 12, 16, 16))
    ozet = defaultdict(lambda: {"istek": 0, "4": 0, "5": 0, "404": set(), "fail": 0, "ok": 0, "ilk": None, "son": None})
    for ol in olaylar:
        adres = ol.ip
        if ol.tur == "auth":
            m = BASARISIZ.search(ol.mesaj) or BASARILI.search(ol.mesaj)
            if not m:
                continue
            adres = m.group("ip")
        if not adres:
            continue
        x = ozet[adres]
        x["ilk"] = x["ilk"] or ol.zaman
        x["son"] = ol.zaman
        if ol.tur == "erisim":
            x["istek"] += 1
            x["4"] += 400 <= ol.kod < 500
            x["5"] += ol.kod >= 500
            if ol.kod == 404:
                x["404"].add(ol.yol)
        elif BASARILI.search(ol.mesaj):
            x["ok"] += 1
        else:
            x["fail"] += 1
    for a, x in sorted(ozet.items(), key=lambda i: -(i[1]["4"] + i[1]["5"] + i[1]["fail"] * 2 + len(i[1]["404"])))[:200]:
        ip.append([a, x["istek"], x["4"], x["5"], len(x["404"]), x["fail"], x["ok"], x["ilk"], x["son"]])
        for c in (8, 9):
            ip.cell(ip.max_row, c).number_format = "DD.MM HH:MM"
    ip.auto_filter.ref = ip.dimensions

    h = wb.create_sheet("Hata İmzaları")
    _baslik(h, ["İmza (normalleştirilmiş)", "Adet", "İlk", "Son", "Örnek"], (70, 8, 16, 16, 90))
    for s, os_ in sorted(imzalar.items(), key=lambda i: -len(i[1])):
        h.append([s, len(os_), os_[0].zaman, os_[-1].zaman, os_[0].mesaj[:300]])
        for c in (3, 4):
            h.cell(h.max_row, c).number_format = "DD.MM HH:MM"
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


# ----------------------------------------------------------------------------
# Akış
# ----------------------------------------------------------------------------

def ayarlar(pencere_dk: int = 10, esik_giris: int = 10, esik_404: int = 20, esik_istek: int = 300, min_artis: int = 10,
            mesai: str = "08-19", yonetim_yolu: str = r"^/(admin|yonetim|wp-admin|manager|panel)\b") -> dict:
    m = re.fullmatch(r"(\d{1,2})-(\d{1,2})", mesai)
    if not m:
        raise SystemExit("--mesai SS-SS biçiminde olmalı, ör. 08-19")
    return {"pencere": timedelta(minutes=pencere_dk), "pencere_dk": pencere_dk, "esik_giris": esik_giris, "esik_404": esik_404,
            "esik_istek": esik_istek, "min_artis": min_artis, "mesai": (int(m.group(1)), int(m.group(2))), "yonetim_yolu": yonetim_yolu}


def calistir(girdiler: list[Path], cikti: Path, yil: int | None = None, **kw) -> dict:
    p = ayarlar(**kw)
    olaylar, ozet = loglari_oku(girdiler, yil or date.today().year)
    bulgular, zaman, imzalar = analiz_et(olaylar, p)
    rapor_yaz(cikti, olaylar, ozet, bulgular, zaman, imzalar, p)
    return {"olaylar": olaylar, "dosyalar": ozet, "bulgular": bulgular, "zaman": zaman, "imzalar": imzalar}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Sunucu ve uygulama loglarında olağandışı hata, oturum ve erişim örüntülerini bulur.")
    p.add_argument("--girdi", type=Path, nargs="+", default=[BURASI / "ornek_veri"], help="Log dosyaları veya klasör (.log, .gz)")
    p.add_argument("--yil", type=int, help="Yılı yazmayan syslog satırları için yıl (varsayılan: bu yıl; örnek veride 2026)")
    p.add_argument("--pencere", type=int, default=10, help="Zaman dilimi ve yoğunluk penceresi, dakika (varsayılan 10)")
    p.add_argument("--esik-giris", type=int, default=10, help="Pencerede bir IP'den başarısız giriş eşiği (varsayılan 10)")
    p.add_argument("--esik-404", type=int, default=20, help="Bir IP'nin farklı 404 yolu eşiği — tarama (varsayılan 20)")
    p.add_argument("--esik-istek", type=int, default=300, help="Bir IP'den dakikada istek eşiği (varsayılan 300)")
    p.add_argument("--min-artis", type=int, default=10, help="Ani artış için dilimde en az olay (varsayılan 10)")
    p.add_argument("--mesai", default="08-19", help="Mesai saatleri SS-SS (varsayılan 08-19; hafta sonu mesai dışı)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "log_anomali_raporu.xlsx")
    a = p.parse_args(argv)
    yil = a.yil or (2026 if a.girdi == [BURASI / "ornek_veri"] else None)
    s = calistir(a.girdi, a.cikti, yil, pencere_dk=a.pencere, esik_giris=a.esik_giris, esik_404=a.esik_404, esik_istek=a.esik_istek,
                 min_artis=a.min_artis, mesai=a.mesai)
    for ad, x in s["dosyalar"].items():
        print(f"[OK] {ad}: {x['cozulen']}/{x['satir']} satır çözüldü")
    say = Counter(b.onem for b in s["bulgular"])
    print(f"[OK] {len(s['bulgular'])} bulgu · " + " · ".join(f"{k} {say[k]}" for k in RENK if say[k]))
    for b in s["bulgular"]:
        if b.onem in ("Kritik", "Yüksek"):
            print(f"[{'X' if b.onem == 'Kritik' else '!'}] {b.tur} · {b.varlik} · {b.aciklama[:110]}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
