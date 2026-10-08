import sys
import tempfile
import unittest
from collections import Counter
from datetime import datetime, timedelta
from pathlib import Path

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))

import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402


class OrnekTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        with tempfile.TemporaryDirectory() as tmp:
            cls.s = main.calistir([KLASOR / "ornek_veri"], Path(tmp) / "l.xlsx", 2026)
            cls.wb = load_workbook(Path(tmp) / "l.xlsx")
        cls.b = {(b.tur, b.varlik): b for b in cls.s["bulgular"]}

    def test_dosyalar(self):
        d = self.s["dosyalar"]
        self.assertEqual({a: (x["tur"], x["cozulen"] == x["satir"]) for a, x in d.items()},
                         {"access.log": ("erisim", True), "app.log": ("uygulama", True), "auth.log": ("auth", True)})

    def test_kritikler(self):
        kritik = {(b.tur, b.varlik) for b in self.s["bulgular"] if b.onem == "Kritik"}
        self.assertEqual(kritik, {("SSH parola deneme", "198.51.100.23"), ("Web parola deneme", "198.51.100.7")})
        self.assertIn("ARDINDAN 07.10 04:15:02 başarılı giriş: root (password)", self.b[("SSH parola deneme", "198.51.100.23")].aciklama)

    def test_tarama_ve_imzalar(self):
        t = self.b[("Dizin / zafiyet taraması", "203.0.113.45")]
        self.assertEqual((t.onem, t.adet), ("Yüksek", 60))
        ssh = self.b[("SSH parola deneme", "203.0.113.45")]
        self.assertIn("25 deneme var olmayan kullanıcı adlarıyla", ssh.aciklama)
        self.assertEqual(self.b[("Saldırı imzası: SQL enjeksiyonu", "203.0.113.99")].onem, "Yüksek")      # 500 döndü
        self.assertEqual(self.b[("Saldırı imzası: Hassas dosya / panel arama", "203.0.113.45")].onem, "Dikkat")

    def test_artislar(self):
        b5 = self.b[("5xx artışı", "-")]
        self.assertEqual((b5.bas, b5.bit, b5.adet), (datetime(2026, 10, 7, 14, 30), datetime(2026, 10, 7, 14, 40), 35))
        yeni = {b.varlik for b in self.s["bulgular"] if b.tur == "Ani ortaya çıkan hata türü"}
        self.assertEqual(len(yeni), 2)                       # db havuzu ve ödeme zaman aşımı; dağınık "stok kaydı" hatası değil
        self.assertTrue(any(v.startswith("[db] Connection pool exhausted") for v in yeni))

    def test_diger(self):
        self.assertIn(("Mesai dışı yönetim erişimi", "192.0.2.10"), self.b)
        self.assertIn(("Yeni IP'den giriş", "muhasebe @ 10.0.0.99"), self.b)
        self.assertIn(("sudo parola hatası", "ahmet"), self.b)
        self.assertNotIn(("Mesai dışı giriş", "deploy @ 10.0.0.5"), self.b)
        self.assertEqual(Counter(b.onem for b in self.s["bulgular"]), Counter({"Kritik": 2, "Yüksek": 7, "Dikkat": 6, "Bilgi": 2}))

    def test_excel(self):
        self.assertEqual(self.wb.sheetnames, ["Özet", "Bulgular", "Zaman Çizelgesi", "IP Özeti", "Hata İmzaları"])
        self.assertEqual(self.wb["Bulgular"]["A2"].value, "Kritik")


class YardimciTest(unittest.TestCase):
    def test_imza(self):
        self.assertEqual(main.imza("Timeout calling https://a.b/x?y=1 after 30000 ms (id='abc')"), "Timeout calling <url:a.b> after <n> ms (id=<s>)")

    def test_en_yogun(self):
        t0 = datetime(2026, 1, 1)
        z = [t0 + timedelta(minutes=m) for m in (0, 1, 2, 30, 31, 32, 33)]
        self.assertEqual(main.en_yogun(z, timedelta(minutes=5)), 4)

    def test_ani_artis(self):
        t0 = datetime(2026, 1, 1)
        dl = [t0 + timedelta(minutes=10 * i) for i in range(20)]
        sayac = Counter({d: 2 for d in dl})
        sayac[dl[12]] = 25
        r, med = main.ani_artis(sayac, dl, timedelta(minutes=10), 10)
        self.assertEqual((len(r), r[0][0], r[0][2], med), (1, dl[12], 25, 2))

    def test_satir_bicimleri(self):
        o = main.satir_coz("2026-10-07T03:00:00.123456+03:00 web01 sshd[1]: Failed password for root from 192.0.2.1 port 1 ssh2", "a", 2026)
        self.assertEqual((o.tur, o.zaman.hour), ("auth", 3))
        o = main.satir_coz('192.0.2.1 - - [07/Oct/2026:10:00:00 +0300] "GET / HTTP/1.1" 200 12', "a", 2026)
        self.assertEqual((o.tur, o.kod, o.ip), ("erisim", 200, "192.0.2.1"))
        self.assertIsNone(main.satir_coz("rastgele metin", "a", 2026))


if __name__ == "__main__":
    unittest.main()
