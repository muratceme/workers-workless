import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
H1, H2 = date(2026, 9, 28), date(2026, 10, 5)


def s(hafta, urun, ciro, adet=0, siparis=0, ziyaret=0, iade=0, kanal="Web", kategori="Giyim"):
    return {"hafta": hafta, "kanal": kanal, "kategori": kategori, "urun": urun, "ciro": D(ciro), "adet": D(adet), "siparis": D(siparis), "ziyaret": D(ziyaret), "iade": D(iade)}


class RaporTesti(unittest.TestCase):
    def test_gostergeler(self):
        r = main.raporla([s(H2, "T", 36000, 120, 100, 2000, 6), s(H1, "T", 30000, 100, 90, 1800, 5)])
        u = r["urunler"][0]
        self.assertEqual((u["donusum"], u["sepet"], u["iade_orani"], u["onceki"], u["degisim"], u["pay"], u["isaret"]), (D("5.00"), D("360.00"), D("5.00"), D(30000), D("20.00"), D("100.00"), []))
        self.assertEqual((r["hafta"], r["onceki_var"]), ((date(2026, 10, 5), date(2026, 10, 11)), True))
        u = main.raporla([s(H2, "T", 100)])["urunler"][0]
        self.assertEqual((u["donusum"], u["sepet"], u["iade_orani"], u["degisim"]), (None, None, None, None))
        self.assertEqual(main.pazartesi(date(2026, 10, 11)), H2)

    def test_isaretler_ve_hafta_secimi(self):
        self.assertEqual([main.raporla([s(H2, "T", 1, 100, iade=i)])["urunler"][0]["isaret"] for i in (15, 16)], [[], ["İade yüksek"]])                        # tam %15 işaretlenmez
        self.assertEqual([main.raporla([s(H2, "T", c), s(H1, "T", 100)])["urunler"][0]["isaret"] for c in (71, 70)], [[], ["Ciro düşüyor"]])                   # tam %30 düşüş işaretlenir
        r = main.raporla([s(H2, "A", 100), s(H1, "B", 50), s(H1, "A", 100)])
        self.assertEqual([(u["urun"], u["ciro"], u["isaret"]) for u in r["urunler"]], [("A", D(100), []), ("B", D(0), ["Bu hafta satış yok"])])
        self.assertEqual((r["genel"]["onceki"], r["genel"]["degisim"], r["isaretli"]), (D(150), D("-33.33"), 1))
        r = main.raporla([s(H2, "A", 100), s(H1, "A", 80)], date(2026, 10, 1))                                    # perşembe verilir, haftası alınır
        self.assertEqual((r["hafta"][0], r["genel"]["ciro"], r["onceki_var"]), (H1, D(80), False))
        self.assertEqual(main.raporla([s(H2, "T", 10, 100, iade=30)], iade_esik=D(30))["urunler"][0]["isaret"], [])
        for kotu in ([], [s(H1, "A", 1)]):
            with self.assertRaises(ValueError):
                main.raporla(kotu, date(2026, 10, 20) if kotu else None)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "e.xlsx"
            r = main.calistir(ORNEK / "eticaret_satislari.csv", cikti)
            g = r["genel"]
            self.assertEqual((g["ciro"], g["onceki"], g["degisim"], g["adet"], g["siparis"], g["ziyaret"], g["donusum"], g["sepet"], g["iade_orani"]),
                             (D(135000), D(134250), D("0.56"), D(410), D(370), D(8800), D("4.20"), D("364.86"), D("9.02")))
            self.assertEqual([(k["kanal"], k["ciro"], k["degisim"], k["donusum"], k["iade_orani"], k["pay"]) for k in r["kanallar"]],
                             [("Pazaryeri A", D(72000), D("21.52"), D("4.14"), D("8.46"), D("53.33")), ("Web Sitesi", D(63000), D("-16.00"), D("4.33"), D("10.00"), D("46.67"))])
            self.assertEqual([(k["kategori"], k["ciro"], k["degisim"]) for k in r["kategoriler"]], [("Giyim", D(129000), D("0.00")), ("Aksesuar", D(6000), D("14.29"))])
            self.assertEqual([(u["kanal"], u["urun"], u["ciro"], u["degisim"], u["iade_orani"], u["isaret"]) for u in r["urunler"]],
                             [("Pazaryeri A", "Basic Tişört", D(66000), D("22.22"), D("10.00"), []), ("Web Sitesi", "Basic Tişört", D(36000), D("20.00"), D("5.00"), []),
                              ("Web Sitesi", "Slim Jean", D(27000), D("-40.00"), D("30.00"), ["İade yüksek", "Ciro düşüyor"]), ("Pazaryeri A", "Şapka", D(6000), D("14.29"), D("0.00"), [])])
            self.assertEqual((r["uyarilar"], load_workbook(cikti).sheetnames), ([], ["Kanal Bazında", "Kategori Bazında", "Ürün Bazında", "Uyarılar"]))

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "s.csv"
            y.write_text("Tarih;Kanal;Ürün;Ziyaret;Sipariş;Satış Adedi;Ciro;İade\n07.10.2026;Web;A;çok;5;6;1.200,00;-1\ndün;Web;B;1;1;1;10\n08.10.2026;;C;;;;yok\n2026-10-09;;D;;;;50\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual([(x["hafta"], x["kanal"], x["urun"], x["ciro"], x["ziyaret"], x["iade"]) for x in l], [(H2, "Web", "A", D("1200.00"), D(0), D(0)), (H2, "Belirtilmemiş", "D", D(50), D(0), D(0))])
            self.assertEqual(sorted(u["tur"] for u in uy), ["Değer okunamadı", "Değer okunamadı", "Tarih veya ciro okunamadı", "Tarih veya ciro okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--satislar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--hafta", "06.10.2026", "--iade-esik", "10", "--dusus-esik", "25"]), 0)
            self.assertEqual(main.main(arg + ["--hafta", "01.01.2026"]), 1)
            self.assertEqual(main.main(arg + ["--hafta", "geçen hafta"]), 2)
            self.assertEqual(main.main(arg + ["--iade-esik", "-5"]), 2)
            self.assertEqual(main.main(["--satislar", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Tarih;Ciro\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
