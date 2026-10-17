"""
E-ticaret Satış Performans Raporu — Gerisi Bende kod bloğu
Kategori ve İçerik Yönetimi › Kategori Yöneticisi

Kanal, kategori ve ürün bazında ciro, adet, dönüşüm ve iade oranını haftalık raporlar.
  Hafta           = pazartesi-pazar; rapor haftası --hafta ile verilen günün haftası, verilmezse verideki son hafta
  Dönüşüm %       = sipariş sayısı ÷ ziyaret × 100
  Ortalama sepet  = ciro ÷ sipariş sayısı
  İade oranı %    = iade adedi ÷ satış adedi × 100       (aynı haftanın kayıtlarından; iadeler çoğunlukla önceki haftaların satışına aittir)
  Ciro değişimi % = (bu hafta ciro − önceki hafta ciro) ÷ önceki hafta ciro × 100
  Pay %           = ciro ÷ haftanın toplam cirosu × 100
  İşaretler       = İade yüksek (iade oranı > --iade-esik 15) · Ciro düşüyor (değişim ≤ −--dusus-esik 30) · Bu hafta satış yok (önceki hafta vardı)
Eşikler varsayılandır; kategoriye göre ayarlayın. Veri panellerden sizin indirdiğiniz dosyadır; kod hiçbir yere bağlanmaz.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"tarih": ("tarih", "gun", "siparis tarihi", "hafta baslangici"), "kanal": ("kanal", "satis kanali", "pazaryeri"), "kategori": ("kategori",), "urun": ("urun", "urun adi", "sku", "stok kodu"),
         "ziyaret": ("ziyaret", "oturum", "goruntulenme", "urun goruntuleme", "ziyaretci"), "siparis": ("siparis", "siparis sayisi", "siparis adedi"), "adet": ("satis adedi", "adet", "satilan adet"),
         "ciro": ("ciro", "net satis", "satis tutari", "tutar"), "iade": ("iade adedi", "iade", "iade adet")}
SAYI = ("ziyaret", "siparis", "adet", "ciro", "iade")
Z = Decimal(0)


def pazartesi(d: date) -> date:
    return d - timedelta(days=d.weekday())


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("tarih", "ciro")):
        if not any(metin(r[a]) for a in SUTUN):
            continue
        kim = f"{metin(r['urun']) or metin(r['kanal']) or 'satır'} (satır {r['_satir']})"
        t, ciro = tarih(r["tarih"]), para(r["ciro"])
        if t is None or ciro is None:
            uy.append(uyari("Yüksek", "Tarih veya ciro okunamadı", kim, f"Tarih '{metin(r['tarih'])}', ciro '{metin(r['ciro'])}'. Satır alınmadı."))
            continue
        x = {"hafta": pazartesi(t), "kanal": metin(r["kanal"]) or "Belirtilmemiş", "kategori": metin(r["kategori"]) or "Belirtilmemiş", "urun": metin(r["urun"]) or "Belirtilmemiş", "ciro": ciro}
        for a in ("ziyaret", "siparis", "adet", "iade"):
            v = para(r[a])
            if v is None and metin(r[a]) or (v is not None and v < 0):
                uy.append(uyari("Orta", "Değer okunamadı", kim, f"{SUTUN[a][0]} '{metin(r[a])}'. 0 sayıldı."))
                v = None
            x[a] = v or Z
        l.append(x)
    return l, uy


def olcut(x: dict, onceki: Decimal | None, toplam: Decimal) -> dict:
    return {"donusum": yuvarla(x["siparis"] / x["ziyaret"] * 100) if x["ziyaret"] else None, "sepet": yuvarla(x["ciro"] / x["siparis"]) if x["siparis"] else None,
            "iade_orani": yuvarla(x["iade"] / x["adet"] * 100) if x["adet"] else None, "onceki": onceki, "degisim": yuvarla((x["ciro"] - onceki) / onceki * 100) if onceki else None,
            "pay": yuvarla(x["ciro"] / toplam * 100) if toplam else Z}


def raporla(satirlar: list[dict], hafta: date | None = None, iade_esik: Decimal = Decimal(15), dusus_esik: Decimal = Decimal(30)) -> dict:
    if not satirlar:
        raise ValueError("Raporlanacak satış satırı yok.")
    hafta = pazartesi(hafta) if hafta else max(s["hafta"] for s in satirlar)
    bu, once = [s for s in satirlar if s["hafta"] == hafta], [s for s in satirlar if s["hafta"] == hafta - timedelta(days=7)]
    if not bu:
        raise ValueError(f"{hafta:%d.%m.%Y} haftasına ait satır yok.")
    toplam = sum((s["ciro"] for s in bu), Z)

    def grupla(alanlar: tuple[str, ...]) -> list[dict]:
        g, o = {}, defaultdict(lambda: Z)
        for s in once:
            o[tuple(katla(s[a]) for a in alanlar)] += s["ciro"]
        for s in bu:
            x = g.setdefault(tuple(katla(s[a]) for a in alanlar), {a: s[a] for a in alanlar} | {a: Z for a in SAYI})
            for a in SAYI:
                x[a] += s[a]
        for k in o:                                    # önceki hafta satılıp bu hafta hiç kaydı olmayanlar
            if k not in g:
                s = next(s for s in once if tuple(katla(s[a]) for a in alanlar) == k)
                g[k] = {a: s[a] for a in alanlar} | {a: Z for a in SAYI}
        for k, x in g.items():
            x |= olcut(x, o.get(k) if once else None, toplam)
        return sorted(g.values(), key=lambda x: (-x["ciro"], *[x[a] for a in alanlar]))

    urunler = grupla(("kanal", "kategori", "urun"))
    for u in urunler:
        i = []
        if u["ciro"] == 0 and u["onceki"]:
            i.append("Bu hafta satış yok")
        if u["iade_orani"] is not None and u["iade_orani"] > iade_esik:
            i.append("İade yüksek")
        if u["degisim"] is not None and u["degisim"] <= -dusus_esik and u["ciro"] > 0:
            i.append("Ciro düşüyor")
        u["isaret"] = i
    genel = grupla(())[0]
    return {"hafta": (hafta, hafta + timedelta(days=6)), "onceki_var": bool(once), "genel": genel, "kanallar": grupla(("kanal",)), "kategoriler": grupla(("kategori",)), "urunler": urunler,
            "esik": (iade_esik, dusus_esik), "isaretli": sum(bool(u["isaret"]) for u in urunler)}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    b, s = r["hafta"]
    ie, de = r["esik"]
    bas = ["Ciro", "Önceki Hafta Ciro", "Ciro Değişimi %", "Ciro Payı %", "Satış Adedi", "Sipariş", "Ziyaret", "Dönüşüm %", "Ortalama Sepet", "İade Adedi", "İade Oranı %"]
    deger = lambda x: [x["ciro"], x["onceki"], x["degisim"], x["pay"], x["adet"], x["siparis"], x["ziyaret"], x["donusum"], x["sepet"], x["iade"], x["iade_orani"]]   # noqa: E731
    gen = (16, 16, 10, 9, 10, 9, 10, 10, 12, 9, 10)

    def sayfa(ws, alanlar, basliklar, liste, toplam=True, isaret=False):
        n = len(alanlar)
        _baslik(ws, [*basliklar, *bas, *(["İşaretler", "Aksiyon (siz doldurun)"] if isaret else [])], (*[24] * n, *gen, 30, 30))
        ps = tuple(n + j for j in (1, 2, 3, 4, 8, 9, 11))
        for x in liste:
            d = x.get("degisim")
            _satir(ws, [*[x[a] for a in alanlar], *deger(x), *([", ".join(x["isaret"]), ""] if isaret else [])], para_s=ps,
                   dolgu={n + 3: (KIRMIZI if d < 0 else YESIL) if d else None} | ({n + 12: KIRMIZI if x["isaret"] else None} if isaret else {}), sar=(n + 12,) if isaret else ())
        if toplam:
            _satir(ws, ["Toplam", *[""] * (n - 1), *deger(r["genel"])], para_s=ps)
            for c in ws[ws.max_row]:
                c.font = Font(bold=True)

    ws = wb.active
    ws.title = "Kanal Bazında"
    sayfa(ws, ("kanal",), ["Kanal"], r["kanallar"])
    ws.append([])
    ws.append([f"Rapor haftası {b:%d.%m.%Y} – {s:%d.%m.%Y}. " + ("" if r["onceki_var"] else "Önceki haftanın verisi yok; değişim hesaplanmadı. ")
               + f"İşaret eşikleri: iade oranı > %{ie.normalize():f}, ciro düşüşü ≥ %{de.normalize():f}. İade oranı aynı haftanın iade ve satış adetlerindendir."])
    sayfa(wb.create_sheet("Kategori Bazında"), ("kategori",), ["Kategori"], r["kategoriler"])
    sayfa(wb.create_sheet("Ürün Bazında"), ("kanal", "kategori", "urun"), ["Kanal", "Kategori", "Ürün"], r["urunler"], toplam=False, isaret=True)
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, hafta: date | None = None, **esik) -> dict:
    satirlar, uy = oku(yol)
    r = raporla(satirlar, hafta, **esik)
    if not r["onceki_var"]:
        uy.append(uyari("Bilgi", "Önceki hafta verisi yok", yol.name, "Ciro değişimi hesaplanmadı. Karşılaştırma için dosyaya önceki haftanın satırlarını da ekleyin."))
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Kanal, kategori ve ürün bazında haftalık e-ticaret satış raporu üretir.")
    p.add_argument("--satislar", type=Path, default=ORNEK / "eticaret_satislari.csv", help="Tarih, Kanal, Kategori, Ürün, Ziyaret, Sipariş, Satış Adedi, Ciro, İade Adedi")
    p.add_argument("--hafta", help="Rapor haftasındaki herhangi bir gün, GG.AA.YYYY (varsayılan: verideki son hafta)")
    p.add_argument("--iade-esik", default="15", help="Bu iade oranının (%%) üstü 'İade yüksek' (varsayılan 15)")
    p.add_argument("--dusus-esik", default="30", help="Önceki haftaya göre bu %% ve üzeri düşüş 'Ciro düşüyor' (varsayılan 30)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "eticaret_satis_raporu.xlsx")
    a = p.parse_args(argv)
    hafta, ie, de = tarih(a.hafta) if a.hafta else None, para(a.iade_esik), para(a.dusus_esik)
    if a.hafta and hafta is None or ie is None or de is None or ie < 0 or de < 0:
        print("[X] --hafta GG.AA.YYYY biçiminde; --iade-esik ve --dusus-esik sıfır veya pozitif sayı olmalı")
        return 2
    if not a.satislar.exists():
        print(f"[X] Dosya bulunamadı: {a.satislar}")
        return 1
    try:
        r = calistir(a.satislar, a.cikti, hafta, iade_esik=ie, dusus_esik=de)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    g, (b, s) = r["genel"], r["hafta"]
    print(f"[OK] {b:%d.%m.%Y} – {s:%d.%m.%Y} · ciro {g['ciro']:,.2f}" + (f" (%{g['degisim']:+})" if g["degisim"] is not None else "") + f" · {g['adet']:f} adet · dönüşüm %{g['donusum']} · iade %{g['iade_orani']}")
    for u in r["urunler"]:
        if u["isaret"]:
            print(f"[!] {u['kanal']} · {u['urun']}: {', '.join(u['isaret'])}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
