"""
Şantiye Puantajı — Gerisi Bende kod bloğu
Şantiye › Saha Mühendisi
Taşeron ve kendi personelinin günlük puantajını kişi × gün cetvelinde ve firma, ekip, imalat bazında toplar.
  Çalışma saati = Saat sütunu; yoksa Gün sütunu (1 / 0,5) × --gun-saat (7,5)
  Adam-gün      = çalışma saati ÷ --gun-saat
  Kontroller    = aynı kişinin bir günde toplam --gunluk-sinir (11) saati aşması · bir haftada (Pzt-Paz) --haftalik (45) saati aşması ·
                  aynı kişinin aynı gün iki firmada yazılması
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"tarih": ("tarih", "gun tarihi"), "personel": ("personel", "ad soyad", "isci", "calisan", "personel adi"), "tc": ("sicil", "sicil no", "personel no", "kimlik no"),
         "firma": ("firma", "taseron", "isveren", "firma adi"), "ekip": ("ekip", "ekip adi", "grup"), "imalat": ("imalat", "is kalemi", "pozisyon is", "yapilan is"),
         "saat": ("saat", "calisma saati", "normal saat"), "gun": ("gun", "puantaj", "yevmiye"), "fm": ("fazla mesai", "fazla calisma", "mesai")}
SIFIR = Decimal(0)
GUN_AD = ("Pzt", "Sal", "Çar", "Per", "Cum", "Cmt", "Paz")


def oku(yol: Path, gun_saat: Decimal = Decimal("7.5")) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("tarih", "personel")):
        ad = metin(r["personel"])
        if not ad:
            continue
        t = tarih(r["tarih"])
        saat, gun, fm = (para(r[a]) if metin(r[a]) else None for a in ("saat", "gun", "fm"))
        bozuk = any(metin(r[a]) and (v is None or v < 0) for a, v in (("saat", saat), ("gun", gun), ("fm", fm)))
        if t is None or bozuk or (saat is None and gun is None):
            uy.append(uyari("Yüksek", "Satır okunamadı", ad, f"Satır {r['_satir']}: tarih, saat ya da gün okunamadı. Satır alınmadı."))
            continue
        s = (saat if saat is not None else gun * gun_saat) + (fm or SIFIR)
        l.append({"tarih": t, "personel": ad, "sicil": metin(r["tc"]), "firma": metin(r["firma"]) or "Kendi personeli", "ekip": metin(r["ekip"]) or "-", "imalat": metin(r["imalat"]) or "-",
                  "saat": s, "fm": fm or SIFIR})
    return l, uy


def puantaj(kayit: list[dict], gun_saat: Decimal = Decimal("7.5"), gunluk: Decimal = Decimal(11), haftalik: Decimal = Decimal(45)) -> dict:
    if not kayit:
        raise ValueError("Puantaj kaydı yok.")
    uy = []
    kisi = {}
    for k in kayit:
        anahtar = katla(k["sicil"]) or katla(k["personel"])
        kisi.setdefault(anahtar, []).append(k)
    gunler = sorted({k["tarih"] for k in kayit})
    cetvel = []
    for anahtar, l in sorted(kisi.items(), key=lambda x: (katla(x[1][0]["firma"]), katla(x[1][0]["personel"]))):
        ad = l[0]["personel"]
        g = defaultdict(lambda: SIFIR)
        firmalar_gun = defaultdict(set)
        hafta = defaultdict(lambda: SIFIR)
        for k in l:
            g[k["tarih"]] += k["saat"]
            firmalar_gun[k["tarih"]].add(k["firma"])
            hafta[k["tarih"].isocalendar()[:2]] += k["saat"]
        for t in sorted(g):
            if g[t] > gunluk:
                uy.append(uyari("Yüksek", "Günlük çalışma sınırı aşıldı", ad, f"{t:%d.%m.%Y} toplam {g[t]:,.2f} saat (sınır {gunluk})."))
            if len(firmalar_gun[t]) > 1:
                uy.append(uyari("Orta", "Aynı gün iki firmada", ad, f"{t:%d.%m.%Y}: {', '.join(sorted(firmalar_gun[t]))}."))
        for (yil, hf), s in sorted(hafta.items()):
            if s > haftalik:
                uy.append(uyari("Bilgi", "Haftalık süre aşıldı", ad, f"{yil} yılı {hf}. hafta toplam {s:,.2f} saat (sınır {haftalik}); aşan kısım fazla çalışma olarak değerlendirilmeli."))
        toplam = sum(g.values(), SIFIR)
        cetvel.append({"personel": ad, "sicil": l[0]["sicil"], "firma": ", ".join(sorted({k["firma"] for k in l}, key=katla)), "ekip": ", ".join(sorted({k["ekip"] for k in l}, key=katla)),
                       "gunler": dict(g), "saat": toplam, "fm": sum((k["fm"] for k in l), SIFIR), "adam_gun": yuvarla(toplam / gun_saat), "calisilan": len(g)})

    def ozet(alan):
        d = {}
        for k in kayit:
            x = d.setdefault(katla(k[alan]), {"ad": k[alan], "saat": SIFIR, "kisi": set(), "gun": set()})
            x["saat"] += k["saat"]
            x["kisi"].add(katla(k["sicil"]) or katla(k["personel"]))
            x["gun"].add(k["tarih"])
        return sorted(({"ad": x["ad"], "saat": x["saat"], "adam_gun": yuvarla(x["saat"] / gun_saat), "kisi": len(x["kisi"]), "gun": len(x["gun"])} for x in d.values()),
                      key=lambda x: (-x["saat"], katla(x["ad"])))
    gunluk_ozet = [{"tarih": t, "kisi": len({katla(k["sicil"]) or katla(k["personel"]) for k in kayit if k["tarih"] == t}), "saat": sum((k["saat"] for k in kayit if k["tarih"] == t), SIFIR)} for t in gunler]
    for x in gunluk_ozet:
        x["adam_gun"] = yuvarla(x["saat"] / gun_saat)
    toplam = sum((k["saat"] for k in kayit), SIFIR)
    return {"cetvel": cetvel, "gunler": gunler, "firma": ozet("firma"), "ekip": ozet("ekip"), "imalat": ozet("imalat"), "gunluk": gunluk_ozet, "saat": toplam,
            "adam_gun": yuvarla(toplam / gun_saat), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, gun_saat: Decimal, gunluk: Decimal) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Puantaj Cetveli"
    g = r["gunler"]
    _baslik(ws, ["Personel", "Sicil", "Firma", "Ekip"] + [f"{t:%d.%m} {GUN_AD[t.weekday()]}" for t in g] + ["Çalışılan Gün", "Toplam Saat", "Fazla Mesai", "Adam-Gün"],
            (22, 10, 18, 14) + (8,) * len(g) + (9, 10, 9, 9), "E2")
    for c in r["cetvel"]:
        degerler = [c["gunler"].get(t) for t in g]
        n = _satir(ws, [c["personel"], c["sicil"], c["firma"], c["ekip"]] + degerler + [c["calisilan"], c["saat"], c["fm"] or None, c["adam_gun"]], para_s=[len(g) + 6, len(g) + 8])
        for i, v in enumerate(degerler):
            if v is not None and v > gunluk:
                ws.cell(n, 5 + i).fill = PatternFill("solid", fgColor=KIRMIZI)
            elif g[i].weekday() == 6 and v:
                ws.cell(n, 5 + i).fill = PatternFill("solid", fgColor=SARI)
    k = _satir(ws, ["TOPLAM", "", "", ""] + [next((x["saat"] for x in r["gunluk"] if x["tarih"] == t), None) for t in g] + [None, r["saat"], None, r["adam_gun"]], para_s=[len(g) + 8])
    for h in ws[k]:
        h.font = Font(bold=True)
    ws.append([])
    ws.append([f"Hücreler çalışma saatidir. Adam-gün = saat ÷ {gun_saat}. Kırmızı: günlük {gunluk} saat üstü; sarı: pazar günü çalışma."])
    ws = wb.create_sheet("Özet")
    _baslik(ws, ["Kırılım", "Ad", "Kişi", "Gün", "Saat", "Adam-Gün"], (10, 26, 7, 7, 10, 10))
    for ad, l in (("Firma", r["firma"]), ("Ekip", r["ekip"]), ("İmalat", r["imalat"])):
        for x in l:
            _satir(ws, [ad, x["ad"], x["kisi"], x["gun"], x["saat"], x["adam_gun"]], para_s=(6,))
    ws = wb.create_sheet("Günlük")
    _baslik(ws, ["Tarih", "Gün", "Kişi", "Saat", "Adam-Gün"], (11, 6, 7, 10, 10))
    for x in r["gunluk"]:
        _satir(ws, [x["tarih"], GUN_AD[x["tarih"].weekday()], x["kisi"], x["saat"], x["adam_gun"]], para_s=(5,), tarih_s=(1,))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, gun_saat: Decimal = Decimal("7.5"), gunluk: Decimal = Decimal(11), haftalik: Decimal = Decimal(45)) -> dict:
    kayit, uy = oku(yol, gun_saat)
    r = puantaj(kayit, gun_saat, gunluk, haftalik)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, gun_saat, gunluk)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Taşeron ve kendi personelinin günlük puantajını ekip ve imalat bazında toplar.")
    p.add_argument("--puantaj", type=Path, help="Tarih, Personel, Sicil, Firma, Ekip, İmalat, Saat (ya da Gün), Fazla Mesai")
    p.add_argument("--gun-saat", default="7,5", help="Bir adam-günün saati (varsayılan 7,5 = 45 saat ÷ 6 gün)")
    p.add_argument("--gunluk-sinir", default="11", help="Günlük toplam çalışma bu saati aşarsa uyarılır (varsayılan 11)")
    p.add_argument("--haftalik", default="45", help="Haftalık toplam bu saati aşarsa bilgi verilir (varsayılan 45)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "santiye_puantaji.xlsx")
    a = p.parse_args(argv)
    yol = a.puantaj or ORNEK / "gunluk_puantaj.csv"
    gs, gl, hf = para(a.gun_saat), para(a.gunluk_sinir), para(a.haftalik)
    if None in (gs, gl, hf) or gs <= 0 or gl <= 0 or hf <= 0:
        print("[X] --gun-saat, --gunluk-sinir ve --haftalik sıfırdan büyük olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, gs, gl, hf)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(r['cetvel'])} kişi · {len(r['gunler'])} gün · {r['saat']:,.2f} saat · {r['adam_gun']:,.2f} adam-gün")
    for f in r["firma"]:
        print(f"     {f['ad']}: {f['kisi']} kişi · {f['adam_gun']:,.2f} adam-gün")
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
