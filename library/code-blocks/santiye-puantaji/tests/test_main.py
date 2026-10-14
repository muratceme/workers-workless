import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def k(gun, ad, firma, imalat, saat, ekip="-"):
    return {"tarih": date(2026, 10, gun), "personel": ad, "sicil": "", "firma": firma, "ekip": ekip, "imalat": imalat, "saat": D(saat), "fm": D(0)}


def veri():
    p1 = [k(g, "P1", "Kendi", "Kalıp", 12 if g == 6 else 8, "A") for g in range(5, 11)]
    return p1 + [k(5, "P2", "Taşeron X", "Demir", "7.5"), k(6, "P2", "Taşeron X", "Demir", "3.75"), k(5, "P3", "Kendi", "Beton", 4), k(5, "p3", "Taşeron X", "Beton", 4)]


class PuantajTesti(unittest.TestCase):
    def test_toplamlar(self):
        r = main.puantaj(veri())
        self.assertEqual([(c["personel"], c["firma"], c["calisilan"], c["saat"], c["adam_gun"]) for c in r["cetvel"]],
                         [("P1", "Kendi", 6, D(52), D("6.93")), ("P3", "Kendi, Taşeron X", 1, D(8), D("1.07")), ("P2", "Taşeron X", 2, D("11.25"), D("1.50"))])
        self.assertEqual([(f["ad"], f["kisi"], f["gun"], f["saat"], f["adam_gun"]) for f in r["firma"]], [("Kendi", 2, 6, D(56), D("7.47")), ("Taşeron X", 2, 2, D("15.25"), D("2.03"))])
        self.assertEqual([(f["ad"], f["saat"]) for f in r["imalat"]], [("Kalıp", D(52)), ("Demir", D("11.25")), ("Beton", D(8))])
        self.assertEqual(([(x["tarih"].day, x["kisi"], x["saat"]) for x in r["gunluk"]][:2], r["saat"], r["adam_gun"]), ([(5, 3, D("23.5")), (6, 2, D("15.75"))], D("71.25"), D("9.50")))
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Bilgi", "Haftalık süre aşıldı", "P1"), ("Orta", "Aynı gün iki firmada", "P3"), ("Yüksek", "Günlük çalışma sınırı aşıldı", "P1")])

    def test_secenekler(self):
        r = main.puantaj(veri(), D(9), D(12), D(60))
        self.assertEqual(([u["tur"] for u in r["uyarilar"]], r["adam_gun"]), (["Aynı gün iki firmada"], D("7.92")))
        with self.assertRaises(ValueError):
            main.puantaj([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "gunluk_puantaj.csv", cikti)
            self.assertEqual(([(f["ad"], f["kisi"], f["adam_gun"]) for f in r["firma"]], r["adam_gun"]),
                             ([("Kendi personeli", 6, D("66.87")), ("Örnek Demir Taşeronu", 6, D("53.80")), ("Örnek Elektrik Taşeronu", 2, D("21.67"))], D("142.33")))
            self.assertEqual([(u["tur"], u["kim"]) for u in r["uyarilar"]], [("Günlük çalışma sınırı aşıldı", "Kemal A."), ("Aynı gün iki firmada", "Cem O."), ("Haftalık süre aşıldı", "Ali R."),
                                                                          ("Haftalık süre aşıldı", "Kemal A."), ("Haftalık süre aşıldı", "Yusuf Ç.")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Puantaj Cetveli", "Özet", "Günlük", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "p.csv"
            y.write_text("Tarih;Ad Soyad;Taşeron;Gün;Fazla Mesai\n05.10.2026;A;;1;2\n05.10.2026;B;X;0,5;\nx;C;;1;\n05.10.2026;D;;;\n05.10.2026;E;;-1;\n05.10.2026;;;1;\n", encoding="utf-8")
            l, uy = main.oku(y, D("7.5"))
            self.assertEqual(([(x["personel"], x["firma"], x["saat"], x["fm"]) for x in l], len(uy)), ([("A", "Kendi personeli", D("9.5"), D(2)), ("B", "X", D("3.75"), D(0))], 3))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--puantaj", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--gun-saat", "9", "--gunluk-sinir", "12", "--haftalik", "50"]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--gun-saat", "0"], ["--gunluk-sinir", "x"], ["--haftalik", "-1"])], [2] * 3)
            self.assertEqual(main.main(["--puantaj", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Tarih;Gün\n05.10.2026;1\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
