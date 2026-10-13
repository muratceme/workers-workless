"""
Üretim Hedef Takibi — Workers / Workless kod bloğu
Satış ve Acente Yönetimi › Bölge Satış Yöneticisi

Branş ve acente bazında prim üretimini aylık ve yıl başından bugüne hedefle karşılaştırır.
  Ay gerçekleşme %        = rapor ayındaki üretim ÷ rapor ayının hedefi × 100
  Kümülatif gerçekleşme % = yıl başından rapor ayına kadar üretim ÷ aynı dönemin hedefi × 100
  Yıllık hedefe göre %    = kümülatif üretim ÷ yıllık hedef × 100
  Kalan hedef             = yıllık hedef − kümülatif üretim;   gereken aylık üretim = kalan hedef ÷ kalan ay sayısı
  Durum                   = Hedefin üstünde (kümülatif ≥ %100) · Hedefe yakın (≥ --esik, %80) · Hedefin gerisinde · Hedefsiz
  Yıllık yazılan hedef    = Ay sütununa yalnız yıl yazılmışsa on iki aya eşit bölünür
Üretim tutarının tanımı (brüt prim, iptal ve zeyiller düşülmüş mü) sizin verinize göredir.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

H_SUTUN = {"acente": ("acente", "acente adi", "satis kanali"), "brans": ("brans", "urun"), "ay": ("ay", "donem", "yil"), "hedef": ("hedef", "hedef prim", "uretim hedefi")}
U_SUTUN = {"acente": ("acente", "acente adi", "satis kanali"), "brans": ("brans", "urun"), "ay": ("tarih", "tanzim tarihi", "ay", "donem"), "prim": ("prim", "uretim", "yazilan prim", "brut prim")}
AYLAR = ("ocak", "subat", "mart", "nisan", "mayis", "haziran", "temmuz", "agustos", "eylul", "ekim", "kasim", "aralik")
AY_AD = ("Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık")
DURUM_RENK = {"Hedefin üstünde": YESIL, "Hedefe yakın": SARI, "Hedefin gerisinde": KIRMIZI, "Hedefsiz": None}


def ay_oku(x) -> tuple[int, int | None] | None:
    """Tarih, 2026-09, 09.2026, Eylül 2026 → (yıl, ay);  yalnız 2026 → (yıl, None)."""
    if isinstance(x, (date, datetime)):
        return x.year, x.month
    s = metin(x)
    if isinstance(x, (int, float)) or re.fullmatch(r"\d{4}(\.0)?", s):
        y = int(float(s))
        return (y, None) if 1990 <= y <= 2100 else None
    if t := tarih(s):
        return t.year, t.month
    if m := re.fullmatch(r"(\d{4})[-./](\d{1,2})", s) or re.fullmatch(r"(\d{1,2})[-./](\d{4})", s):
        a, b = int(m.group(1)), int(m.group(2))
        y, ay = (a, b) if a > 31 else (b, a)
        return (y, ay) if 1 <= ay <= 12 else None
    k = katla(s).split()
    if len(k) == 2 and k[0] in AYLAR and k[1].isdigit():
        return int(k[1]), AYLAR.index(k[0]) + 1
    return None


def oku(yol: Path, sozluk: dict, alan: str, ad: str, yillik: bool) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, sozluk, ("acente", "ay", alan)):
        acente = metin(r["acente"])
        if not acente:
            continue
        ay, tutar = ay_oku(r["ay"]), para(r[alan])
        if ay is None or tutar is None or (ay[1] is None and not yillik) or (alan == "hedef" and tutar < 0):
            uy.append(uyari("Yüksek", f"{ad} satırı okunamadı", acente, f"{yol.name} satır {r['_satir']}: ay '{metin(r['ay'])}' ya da tutar '{metin(r[alan])}' okunamadı. Satır alınmadı."))
            continue
        l.append({"acente": acente, "brans": metin(r["brans"]) or "(branş yok)", "ay": ay, "tutar": tutar})
    return l, uy


def _oran(pay: Decimal, payda: Decimal) -> Decimal | None:
    return yuvarla(pay * 100 / payda) if payda > 0 else None


def takip(hedefler: list[dict], uretim: list[dict], ay: tuple[int, int] | None = None, esik: Decimal = Decimal(80)) -> dict:
    if not hedefler and not uretim:
        raise ValueError("Hedef ya da üretim satırı yok.")
    if ay is None:
        aylar = [u["ay"] for u in uretim] or [h["ay"] for h in hedefler if h["ay"][1]]
        if not aylar:
            raise ValueError("Rapor ayı belirlenemedi; --ay ile verin (ör. 09.2026).")
        ay = max(aylar)
    yil, m = ay
    ad, h, g, disi = {}, defaultdict(lambda: defaultdict(Decimal)), defaultdict(lambda: defaultdict(Decimal)), 0
    for kaynak, hedef_mi in ((hedefler, True), (uretim, False)):
        for s in kaynak:
            if s["ay"][0] != yil:
                disi += 1
                continue
            k = (katla(s["acente"]), katla(s["brans"]))
            ad.setdefault(k, (s["acente"], s["brans"]))
            if s["ay"][1] is None:
                for i in range(1, 13):
                    h[k][i] += s["tutar"] / 12
            else:
                (h if hedef_mi else g)[k][s["ay"][1]] += s["tutar"]
    if not ad:
        raise ValueError(f"{yil} yılına ait hedef ya da üretim yok.")

    def gosterge(anahtarlar: list[tuple]) -> dict:
        t = lambda kaynak, aylar: yuvarla(sum((kaynak[k][i] for k in anahtarlar for i in aylar), Decimal(0)))   # noqa: E731
        ah, ag, kh, kg, yh = t(h, (m,)), t(g, (m,)), t(h, range(1, m + 1)), t(g, range(1, m + 1)), t(h, range(1, 13))
        kalan, oran = max(yh - kg, Decimal(0)), _oran(kg, kh)
        return {"ay_hedef": ah, "ay_gercek": ag, "ay_oran": _oran(ag, ah), "kum_hedef": kh, "kum_gercek": kg, "kum_oran": oran, "fark": kg - kh, "yil_hedef": yh, "yil_oran": _oran(kg, yh), "kalan": kalan,
                "gereken": yuvarla(kalan / (12 - m)) if m < 12 and yh > 0 else None, "durum": "Hedefsiz" if oran is None else "Hedefin üstünde" if oran >= 100 else "Hedefe yakın" if oran >= esik else "Hedefin gerisinde"}

    def kirilim(i: int) -> list[dict]:
        grup = {}
        for k in ad:
            grup.setdefault(k[i], []).append(k)
        return sorted(({"ad": ad[ks[0]][i], **gosterge(ks)} for ks in grup.values()), key=lambda s: (s["kum_oran"] is None, s["kum_oran"] or 0, katla(s["ad"])))

    detay = sorted(({"acente": ad[k][0], "brans": ad[k][1], **gosterge([k])} for k in ad), key=lambda s: (katla(s["acente"]), katla(s["brans"])))
    acenteler, uy = kirilim(0), []
    for a in acenteler:
        if a["durum"] == "Hedefin gerisinde":
            uy.append(uyari("Orta", "Acente hedefin gerisinde", a["ad"], f"Kümülatif üretim {a['kum_gercek']:,.2f}, hedef {a['kum_hedef']:,.2f} (%{a['kum_oran']}; eşik %{esik}). Yıllık hedef için kalan {12 - m} ayda aylık "
                            + (f"{a['gereken']:,.2f}" if a["gereken"] is not None else "—") + " gerekiyor."))
    for d in detay:
        kim = f"{d['acente']} · {d['brans']}"
        if d["kum_hedef"] > 0 and d["kum_gercek"] <= 0:
            uy.append(uyari("Yüksek", "Hedefi olan branşta üretim yok", kim, f"Yıl başından {AY_AD[m - 1]} sonuna kadar hedef {d['kum_hedef']:,.2f}, üretim {d['kum_gercek']:,.2f}."))
        elif d["durum"] == "Hedefsiz" and d["kum_gercek"] != 0:
            uy.append(uyari("Bilgi", "Hedefsiz üretim", kim, f"Kümülatif üretim {d['kum_gercek']:,.2f} ama hedef tanımlı değil."))
    if disi:
        uy.append(uyari("Bilgi", "Rapor yılı dışındaki satırlar", str(yil), f"{disi} satır {yil} yılına ait olmadığı için hesaba girmedi."))
    aylar = [{"ay": AY_AD[i - 1], "hedef": (hh := yuvarla(sum((h[k][i] for k in ad), Decimal(0)))), "gercek": (gg := yuvarla(sum((g[k][i] for k in ad), Decimal(0)))), "oran": _oran(gg, hh), "kapsam": i <= m} for i in range(1, 13)]
    return {"ay": ay, "detay": detay, "acenteler": acenteler, "branslar": kirilim(1), "toplam": gosterge(list(ad)), "aylar": aylar, "uyarilar": uy}


BASLIK = ["Durum", "Ay Hedefi", "Ay Üretimi", "Ay %", "Kümülatif Hedef", "Kümülatif Üretim", "Kümülatif %", "Fark", "Yıllık Hedef", "Yıllık Hedefe Göre %", "Kalan Hedef", "Gereken Aylık Üretim"]


def _deger(s: dict) -> list:
    return [s["durum"], s["ay_hedef"], s["ay_gercek"], s["ay_oran"], s["kum_hedef"], s["kum_gercek"], s["kum_oran"], s["fark"], s["yil_hedef"], s["yil_oran"], s["kalan"], s["gereken"]]


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    yil, m = r["ay"]
    for i, (baslik, alan, ad) in enumerate((("Acenteler", "acenteler", "Acente"), ("Branşlar", "branslar", "Branş"))):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = baslik
        _baslik(ws, [ad] + BASLIK, (30, 18) + (16, 16, 9) * 2 + (16, 16, 11, 16, 16), "B2")
        for s in r[alan]:
            _satir(ws, [s["ad"]] + _deger(s), para_s=range(3, 14), dolgu={2: DURUM_RENK[s["durum"]], 9: KIRMIZI if s["fark"] < 0 else None})
        n = _satir(ws, ["TOPLAM"] + _deger(r["toplam"]), para_s=range(3, 14))
        for h in ws[n]:
            h.font = Font(bold=True)
        ws.append([])
        ws.append([f"Rapor ayı: {AY_AD[m - 1]} {yil}. Kümülatif = Ocak–{AY_AD[m - 1]}."])
    ws = wb.create_sheet("Acente × Branş")
    _baslik(ws, ["Acente", "Branş"] + BASLIK, (30, 18, 18) + (16, 16, 9) * 2 + (16, 16, 11, 16, 16), "C2")
    for s in r["detay"]:
        _satir(ws, [s["acente"], s["brans"]] + _deger(s), para_s=range(4, 15), dolgu={3: DURUM_RENK[s["durum"]], 10: KIRMIZI if s["fark"] < 0 else None})
    ws = wb.create_sheet("Aylar")
    _baslik(ws, ["Ay", "Hedef", "Üretim", "Gerçekleşme %", "Kümülatife Dahil"], (14, 18, 18, 13, 12))
    for a in r["aylar"]:
        _satir(ws, [a["ay"], a["hedef"], a["gercek"], a["oran"], "Evet" if a["kapsam"] else None], para_s=(2, 3, 4), dolgu={4: KIRMIZI if a["kapsam"] and a["oran"] is not None and a["oran"] < 100 else None})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(h_yol: Path, u_yol: Path, cikti: Path, ay: tuple[int, int] | None = None, esik: Decimal = Decimal(80)) -> dict:
    hedefler, uy1 = oku(h_yol, H_SUTUN, "hedef", "Hedef", True)
    uretim, uy2 = oku(u_yol, U_SUTUN, "prim", "Üretim", False)
    r = takip(hedefler, uretim, ay, esik)
    r["uyarilar"] = sirala(uy1 + uy2 + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Branş ve acente bazında prim üretimini aylık ve kümülatif hedefle karşılaştırır.")
    p.add_argument("--hedefler", type=Path, help="Acente, Branş, Ay (ör. 09.2026; yalnız yıl yazılırsa on iki aya bölünür), Hedef")
    p.add_argument("--uretim", type=Path, help="Acente, Branş, Tarih ya da Ay, Prim")
    p.add_argument("--ay", help="Rapor ayı (ör. 09.2026); verilmezse üretim dosyasındaki son ay")
    p.add_argument("--esik", default="80", help="Kümülatif gerçekleşme bu yüzdenin altındaysa 'gerisinde' sayılır (varsayılan 80)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "uretim_hedef_takibi.xlsx")
    a = p.parse_args(argv)
    if (a.hedefler is None) != (a.uretim is None):
        print("[X] --hedefler ve --uretim birlikte verilmeli")
        return 2
    h_yol, u_yol = a.hedefler or ORNEK / "hedefler.csv", a.uretim or ORNEK / "uretim.csv"
    ay, esik = ay_oku(a.ay) if a.ay else None, para(a.esik)
    if (a.ay and (ay is None or ay[1] is None)) or esik is None or not 0 < esik <= 100:
        print("[X] --ay AA.YYYY biçiminde olmalı (ör. 09.2026); --esik 0 ile 100 arasında olmalı")
        return 2
    for yol in (h_yol, u_yol):
        if not yol.exists():
            print(f"[X] Dosya bulunamadı: {yol}")
            return 1
    try:
        r = calistir(h_yol, u_yol, a.cikti, ay, esik)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {AY_AD[r['ay'][1] - 1]} {r['ay'][0]} · {len(r['acenteler'])} acente · ay üretimi {t['ay_gercek']:,.2f}" + (f" (%{t['ay_oran']})" if t["ay_oran"] is not None else "") + f" · kümülatif {t['kum_gercek']:,.2f}"
          + (f" (%{t['kum_oran']})" if t["kum_oran"] is not None else ""))
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
