import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def s(acente, brans, ay, tutar, yil=2026):
    return {"acente": acente, "brans": brans, "ay": (yil, ay), "tutar": D(tutar)}


def veri():
    h = [s("A", "Kasko", m, 100) for m in (1, 2, 3, 4)] + [s("A", "Trafik", m, 50) for m in (1, 2, 3)] + [s("B", "Kasko", m, 200) for m in (1, 2, 3)]
    u = [s("A", "Kasko", 1, 80), s("a", "KASKO", 2, 120), s("A", "Kasko", 3, 130), s("A", "Kasko", 4, 999), s("A", "Trafik", 1, 30), s("A", "Trafik", 3, 30), s("C", "Yangın", 3, 40)]
    return h, u


def k(x):
    return (x["durum"], x["ay_hedef"], x["ay_gercek"], x["ay_oran"], x["kum_hedef"], x["kum_gercek"], x["kum_oran"], x["fark"], x["yil_hedef"], x["yil_oran"], x["kalan"], x["gereken"])


class TakipTesti(unittest.TestCase):
    def test_gostergeler(self):
        r = main.takip(*veri(), ay=(2026, 3))
        self.assertEqual([(x["acente"], x["brans"]) + k(x) for x in r["detay"]],
                         [("A", "Kasko", "Hedefin üstünde", D(100), D(130), D("130.00"), D(300), D(330), D("110.00"), D(30), D(400), D("82.50"), D(70), D("7.78")),
                          ("A", "Trafik", "Hedefin gerisinde", D(50), D(30), D("60.00"), D(150), D(60), D("40.00"), D(-90), D(150), D("40.00"), D(90), D("10.00")),
                          ("B", "Kasko", "Hedefin gerisinde", D(200), D(0), D("0.00"), D(600), D(0), D("0.00"), D(-600), D(600), D("0.00"), D(600), D("66.67")),
                          ("C", "Yangın", "Hedefsiz", D(0), D(40), None, D(0), D(40), None, D(40), D(0), None, D(0), None)])
        self.assertEqual([(x["ad"], x["durum"], x["kum_hedef"], x["kum_gercek"], x["kum_oran"]) for x in r["acenteler"]],
                         [("B", "Hedefin gerisinde", D(600), D(0), D("0.00")), ("A", "Hedefe yakın", D(450), D(390), D("86.67")), ("C", "Hedefsiz", D(0), D(40), None)])
        self.assertEqual([(x["ad"], x["kum_hedef"], x["kum_gercek"], x["kum_oran"]) for x in r["branslar"]], [("Kasko", D(900), D(330), D("36.67")), ("Trafik", D(150), D(60), D("40.00")), ("Yangın", D(0), D(40), None)])
        self.assertEqual(k(r["toplam"]), ("Hedefin gerisinde", D(350), D(200), D("57.14"), D(1050), D(430), D("40.95"), D(-620), D(1150), D("37.39"), D(720), D("80.00")))
        self.assertEqual([(a["ay"], a["hedef"], a["gercek"], a["oran"], a["kapsam"]) for a in r["aylar"][:5]],
                         [("Ocak", D(350), D(110), D("31.43"), True), ("Şubat", D(350), D(120), D("34.29"), True), ("Mart", D(350), D(200), D("57.14"), True), ("Nisan", D(100), D(999), D("999.00"), False), ("Mayıs", D(0), D(0), None, False)])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]), [("Bilgi", "Hedefsiz üretim", "C · Yangın"), ("Orta", "Acente hedefin gerisinde", "B"), ("Yüksek", "Hedefi olan branşta üretim yok", "B · Kasko")])

    def test_yillik_esik_ve_ay(self):
        h = [{"acente": "A", "brans": "K", "ay": (2026, None), "tutar": D(1200)}, s("A", "K", 5, 999, 2025)]
        u = [s("A", "K", m, v) for m, v in ((1, 100), (2, 60), (3, 80))] + [s("A", "K", 12, 500, 2025)]
        r = main.takip(h, u)
        self.assertEqual((r["ay"], k(r["toplam"]), [u["tur"] for u in r["uyarilar"]]),
                         ((2026, 3), ("Hedefe yakın", D("100.00"), D(80), D("80.00"), D("300.00"), D(240), D("80.00"), D("-60.00"), D("1200.00"), D("20.00"), D("960.00"), D("106.67")), ["Rapor yılı dışındaki satırlar"]))
        d = lambda **kw: main.takip(h, u, **kw)["toplam"]   # noqa: E731
        self.assertEqual((d(esik=D(81))["durum"], d(ay=(2026, 1))["durum"], d(ay=(2026, 12))["gereken"], d(ay=(2026, 12))["kum_oran"], d(ay=(2025, 12))["kum_gercek"]), ("Hedefin gerisinde", "Hedefin üstünde", None, D("20.00"), D(500)))
        self.assertEqual(main.takip([s("A", "K", 2, 10)], [])["ay"], (2026, 2))
        for hh, uu, kw in (([], [], {}), ([{"acente": "A", "brans": "K", "ay": (2026, None), "tutar": D(1)}], [], {}), ([s("A", "K", 1, 5)], [], {"ay": (2030, 1)})):
            with self.assertRaises(ValueError):
                main.takip(hh, uu, **kw)

    def test_ay_oku(self):
        self.assertEqual([main.ay_oku(x) for x in ("2026-09", "09.2026", "9/2026", "2026/9", "Eylül 2026", "ŞUBAT 2026", "15.09.2026", date(2026, 9, 3), "2026", 2026, 2026.0, "13.2026", "x", "", "99")],
                         [(2026, 9)] * 5 + [(2026, 2), (2026, 9), (2026, 9), (2026, None), (2026, None), (2026, None), None, None, None, None])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "hedefler.csv", ORNEK / "uretim.csv", cikti)
            self.assertEqual(r["ay"], (2026, 9))
            self.assertEqual([(x["acente"], x["brans"], x["durum"], x["ay_oran"], x["kum_hedef"], x["kum_gercek"], x["kum_oran"]) for x in r["detay"]],
                             [("Deneme Sigorta Aracılık", "Kasko", "Hedefin gerisinde", D("60.00"), D(675000), D(445000), D("65.93")), ("Deneme Sigorta Aracılık", "Yangın", "Hedefin üstünde", D("120.00"), D(270000), D(276000), D("102.22")),
                              ("Numune Acente", "Sağlık", "Hedefsiz", None, D(0), D(15000), None), ("Numune Acente", "Trafik", "Hedefin gerisinde", D("0.00"), D(360000), D(0), D("0.00")),
                              ("Örnek Acente", "Kasko", "Hedefe yakın", D("120.00"), D(900000), D(880000), D("97.78")), ("Örnek Acente", "Trafik", "Hedefin üstünde", D("80.00"), D(450000), D(480000), D("106.67"))])
            self.assertEqual([(x["ad"], x["durum"], x["kum_oran"]) for x in r["acenteler"]],
                             [("Numune Acente", "Hedefin gerisinde", D("4.17")), ("Deneme Sigorta Aracılık", "Hedefin gerisinde", D("76.30")), ("Örnek Acente", "Hedefin üstünde", D("100.74"))])
            self.assertEqual([(x["ad"], x["kum_oran"]) for x in r["branslar"]], [("Trafik", D("59.26")), ("Kasko", D("84.13")), ("Yangın", D("102.22")), ("Sağlık", None)])
            t = r["toplam"]
            self.assertEqual((t["ay_hedef"], t["ay_gercek"], t["ay_oran"], t["kum_hedef"], t["kum_gercek"], t["kum_oran"]), (D(295000), D(256000), D("86.78"), D(2655000), D(2096000), D("78.95")))
            ok = next(x for x in r["detay"] if x["acente"] == "Örnek Acente" and x["brans"] == "Kasko")
            self.assertEqual((ok["kalan"], ok["gereken"]), (D(320000), D("106666.67")))
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Yüksek", "Hedefi olan branşta üretim yok", "Numune Acente · Trafik"), ("Orta", "Acente hedefin gerisinde", "Deneme Sigorta Aracılık"), ("Orta", "Acente hedefin gerisinde", "Numune Acente"),
                              ("Bilgi", "Hedefsiz üretim", "Numune Acente · Sağlık")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Acenteler", "Branşlar", "Acente × Branş", "Aylar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            h, u = Path(tmp) / "h.csv", Path(tmp) / "u.csv"
            h.write_text("Acente Adı;Ürün;Dönem;Hedef Prim\nA;Kasko;2026;12.000\nA;;03.2026;500,50\nB;K;x;5\nC;K;2026;-1\n;K;2026;5\n", encoding="utf-8")
            u.write_text("Acente;Branş;Tanzim Tarihi;Brüt Prim\nA;Kasko;15.03.2026;1.000\nA;Kasko;20.03.2026;-200\nA;Kasko;2026;50\nA;Kasko;03.2026;y\n", encoding="utf-8")
            hl, uy = main.oku(h, main.H_SUTUN, "hedef", "Hedef", True)
            self.assertEqual(([(x["acente"], x["brans"], x["ay"], x["tutar"]) for x in hl], [(x["tur"], x["kim"]) for x in uy]),
                             ([("A", "Kasko", (2026, None), D(12000)), ("A", "(branş yok)", (2026, 3), D("500.50"))], [("Hedef satırı okunamadı", "B"), ("Hedef satırı okunamadı", "C")]))
            ul, uy = main.oku(u, main.U_SUTUN, "prim", "Üretim", False)
            self.assertEqual(([(x["ay"], x["tutar"]) for x in ul], [x["tur"] for x in uy]), ([((2026, 3), D(1000)), ((2026, 3), D(-200))], ["Üretim satırı okunamadı"] * 2))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--hedefler", str(h), "--uretim", str(u), "--cikti", c]
            self.assertEqual(main.main(arg + ["--ay", "03.2026", "--esik", "90"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--ay", "x"], ["--ay", "2026"], ["--esik", "0"], ["--esik", "101"])] + [main.main(["--hedefler", str(h)])], [2] * 5)
            self.assertEqual(main.main(arg + ["--ay", "03.2031"]), 1)
            self.assertEqual(main.main(["--hedefler", str(h), "--uretim", str(Path(tmp) / "yok.csv")]), 1)
            u.write_text("Acente;Branş\nA;B\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
