"""
Eksper Performans Analizi — Gerisi Bende kod bloğu
Hasar Yönetimi › Hasar Müdürü

Eksperleri rapor süresi, tespit tutarı ile ödeme arasındaki fark ve itiraz oranına göre karşılaştırır.
  Rapor süresi (gün)     = rapor tarihi − atama tarihi;   ortalama ve ortanca eksper bazında
  Süre aşımı %           = rapor süresi --sla (15) günü aşan dosya ÷ raporlanan dosya × 100
  Bekleyen dosya         = rapor tarihi boş;   bekleme = durum tarihi − atama tarihi
  Tespit–ödeme farkı %   = (Σ ödenen − Σ tespit) ÷ Σ tespit × 100          (tespit ve ödemesi olan dosyalar)
  Mutlak sapma %         = Σ |ödenen − tespit| ÷ Σ tespit × 100            (artı ve eksi farklar birbirini götürmez)
  İtiraz oranı %         = itiraz edilen dosya ÷ raporlanan dosya × 100
  Değerlendirme          = raporlanan dosyası --min (3) ve üstü olan eksperde gösterge genel değerin (1 + --sapma %20) katını aşıyorsa işaretlenir
Kod eksperin raporunu değerlendirmez; fark ve itirazın nedeni dosya bazında incelenmelidir.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"no": ("dosya no", "hasar dosya no", "hasar no"), "eksper": ("eksper", "eksper adi", "eksper firmasi"), "brans": ("brans", "urun"), "atama": ("atama tarihi", "gorevlendirme tarihi", "tayin tarihi"),
         "rapor": ("rapor tarihi", "rapor teslim tarihi"), "tespit": ("tespit tutari", "eksper tespit tutari", "ekspertiz tutari"), "odenen": ("odenen tutar", "odenen", "odeme tutari"), "itiraz": ("itiraz", "itiraz var mi"),
         "ucret": ("eksper ucreti", "ekspertiz ucreti", "ucret")}
EVET = {"evet", "e", "var", "x", "true", "1"}


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("eksper", "atama")):
        eksper, no = metin(r["eksper"]), metin(r["no"]) or f"Satır {r['_satir']}"
        if not eksper:
            continue
        atama, rapor = tarih(r["atama"]), tarih(r["rapor"])
        t = {a: para(r[a]) if metin(r[a]) else None for a in ("tespit", "odenen", "ucret")}
        if atama is None or (metin(r["rapor"]) and rapor is None) or any(metin(r[a]) and (t[a] is None or t[a] < 0) for a in t):
            uy.append(uyari("Yüksek", "Dosya okunamadı", no, f"Satır {r['_satir']}: atama tarihi, rapor tarihi ya da tutarlardan biri okunamadı. Satır alınmadı."))
            continue
        l.append({"no": no, "eksper": eksper, "brans": metin(r["brans"]), "atama": atama, "rapor": rapor, **t, "itiraz": katla(r["itiraz"]) in EVET})
    return l, uy


def _oran(pay, payda) -> Decimal | None:
    return yuvarla(Decimal(pay) * 100 / Decimal(payda)) if payda else None


def _medyan(d: list[int]) -> Decimal | None:
    s, n = sorted(d), len(d)
    return None if not n else Decimal(s[n // 2]) if n % 2 else yuvarla(Decimal(s[n // 2 - 1] + s[n // 2]) / 2, "0.1")


def gosterge(l: list[dict], sla: int) -> dict:
    rap, bek, cift = [d for d in l if d["sure"] is not None], [d for d in l if d["bekleme"] is not None], [d for d in l if d["fark"] is not None]
    sureler, tespit, odenen = [d["sure"] for d in rap], sum((d["tespit"] for d in cift), Decimal(0)), sum((d["odenen"] for d in cift), Decimal(0))
    ucretler = [d["ucret"] for d in l if d["ucret"] is not None]
    return {"dosya": len(l), "raporlanan": len(rap), "ort_sure": yuvarla(Decimal(sum(sureler)) / len(sureler)) if sureler else None, "medyan": _medyan(sureler), "asim": sum(1 for s in sureler if s > sla),
            "asim_oran": _oran(sum(1 for s in sureler if s > sla), len(rap)), "bekleyen": len(bek), "enb_bekleme": max((d["bekleme"] for d in bek), default=None), "tespit": tespit, "odenen": odenen, "fark": _oran(odenen - tespit, tespit),
            "mutlak": _oran(sum((abs(d["fark"]) for d in cift), Decimal(0)), tespit), "itiraz": sum(1 for d in rap if d["itiraz"]), "itiraz_oran": _oran(sum(1 for d in rap if d["itiraz"]), len(rap)),
            "ucret": sum(ucretler, Decimal(0)), "ort_ucret": yuvarla(sum(ucretler, Decimal(0)) / len(ucretler)) if ucretler else None}


def analiz(dosyalar: list[dict], bugun: date, sla: int = 15, asgari: int = 3, sapma: Decimal = Decimal(20)) -> dict:
    if not dosyalar:
        raise ValueError("Analiz edilecek ekspertiz dosyası yok.")
    uy = []
    for d in dosyalar:
        kim = f"{d['no']} · {d['eksper']}"
        if d["rapor"] and d["rapor"] < d["atama"]:
            uy.append(uyari("Orta", "Rapor tarihi atamadan önce", kim, f"Atama {d['atama']:%d.%m.%Y}, rapor {d['rapor']:%d.%m.%Y}. Süre sıfır sayıldı; tarihleri kontrol edin."))
        d["sure"] = max((d["rapor"] - d["atama"]).days, 0) if d["rapor"] else None
        d["bekleme"] = None if d["rapor"] else max((bugun - d["atama"]).days, 0)
        d["fark"] = d["odenen"] - d["tespit"] if d["rapor"] and d["tespit"] and d["odenen"] is not None else None
        d["fark_oran"] = _oran(d["fark"], d["tespit"]) if d["fark"] is not None else None
        if d["bekleme"] is not None and d["bekleme"] > sla:
            uy.append(uyari("Orta", "Rapor gecikti", kim, f"Atama {d['atama']:%d.%m.%Y}; {d['bekleme']} gündür rapor yok (hedef {sla} gün)."))
    genel, kat = gosterge(dosyalar, sla), (100 + sapma) / 100
    grup = {}
    for d in dosyalar:
        grup.setdefault(katla(d["eksper"]), []).append(d)
    eksperler = []
    for g in grup.values():
        e = {"eksper": g[0]["eksper"], **gosterge(g, sla)}
        isaret = []
        if e["raporlanan"] >= asgari:
            if e["ort_sure"] is not None and genel["ort_sure"] and e["ort_sure"] > genel["ort_sure"] * kat:
                isaret.append("Süre uzun")
            if e["mutlak"] is not None and genel["mutlak"] and e["mutlak"] > genel["mutlak"] * kat:
                isaret.append("Tespit–ödeme farkı yüksek")
            if e["itiraz"] >= 2 and genel["itiraz_oran"] and e["itiraz_oran"] > genel["itiraz_oran"] * kat:
                isaret.append("İtiraz yüksek")
        e["durum"] = "Az veri" if e["raporlanan"] < asgari else "; ".join(isaret) or "Olağan"
        if isaret:
            uy.append(uyari("Orta", "Eksper genel değerlerden sapıyor", e["eksper"], f"{e['durum']}. Ortalama süre {e['ort_sure']} gün (genel {genel['ort_sure']}), mutlak sapma %{e['mutlak']} (genel %{genel['mutlak']}), "
                            f"itiraz %{e['itiraz_oran']} (genel %{genel['itiraz_oran']}); {e['raporlanan']} raporlanan dosya."))
        eksperler.append(e)
    eksperler.sort(key=lambda e: (-e["dosya"], katla(e["eksper"])))
    return {"dosyalar": dosyalar, "eksperler": eksperler, "genel": genel, "bekleyenler": sorted((d for d in dosyalar if d["bekleme"] is not None), key=lambda d: (-d["bekleme"], d["no"])), "uyarilar": uy}


def _deger(e: dict) -> list:
    return [e["dosya"], e["raporlanan"], e["ort_sure"], e["medyan"], e["asim_oran"], e["bekleyen"] or None, e["enb_bekleme"], e["tespit"], e["odenen"], e["fark"], e["mutlak"], e["itiraz"] or None, e["itiraz_oran"], e["ucret"] or None, e["ort_ucret"]]


def rapor_yaz(cikti: Path, r: dict, bugun: date, sla: int) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Eksperler"
    _baslik(ws, ["Eksper", "Değerlendirme", "Dosya", "Raporlanan", "Ort. Rapor Süresi (gün)", "Ortanca Süre (gün)", f"Süre Aşımı % (> {sla} gün)", "Bekleyen Dosya", "En Uzun Bekleme (gün)", "Tespit Tutarı", "Ödenen Tutar", "Fark %",
                 "Mutlak Sapma %", "İtiraz", "İtiraz %", "Toplam Ücret", "Ort. Ücret"], (26, 34, 8, 10, 11, 10, 12, 9, 10, 16, 16, 9, 10, 8, 9, 14, 12), "B2")
    for e in r["eksperler"]:
        _satir(ws, [e["eksper"], e["durum"]] + _deger(e), para_s=(5, 6, 7, 10, 11, 12, 13, 15, 16, 17), dolgu={2: None if e["durum"] == "Az veri" else YESIL if e["durum"] == "Olağan" else SARI}, sar=(2,))
    n = _satir(ws, ["GENEL", None] + _deger(r["genel"]), para_s=(5, 6, 7, 10, 11, 12, 13, 15, 16, 17))
    for h in ws[n]:
        h.font = Font(bold=True)
    ws.append([])
    ws.append([f"Durum tarihi: {bugun:%d.%m.%Y}. Tespit ve ödenen tutarlar yalnız ikisi de yazılmış, raporlanmış dosyaları kapsar."])
    ws = wb.create_sheet("Bekleyen Raporlar")
    _baslik(ws, ["Dosya No", "Eksper", "Branş", "Atama Tarihi", "Bekleme (gün)"], (16, 26, 16, 12, 11))
    for d in r["bekleyenler"]:
        _satir(ws, [d["no"], d["eksper"], d["brans"], d["atama"], d["bekleme"]], tarih_s=(4,), dolgu={5: KIRMIZI if d["bekleme"] > sla else None})
    ws = wb.create_sheet("Dosyalar")
    _baslik(ws, ["Dosya No", "Eksper", "Branş", "Atama", "Rapor", "Süre (gün)", "Tespit Tutarı", "Ödenen Tutar", "Fark", "Fark %", "İtiraz", "Eksper Ücreti"], (16, 26, 16, 12, 12, 9, 16, 16, 15, 9, 8, 13))
    for d in sorted(r["dosyalar"], key=lambda d: (katla(d["eksper"]), d["atama"], d["no"])):
        _satir(ws, [d["no"], d["eksper"], d["brans"], d["atama"], d["rapor"], d["sure"], d["tespit"], d["odenen"], d["fark"], d["fark_oran"], "Evet" if d["itiraz"] else None, d["ucret"]], para_s=(7, 8, 9, 10, 12), tarih_s=(4, 5),
               dolgu={6: SARI if d["sure"] is not None and d["sure"] > sla else None, 11: SARI if d["itiraz"] else None})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, sla: int = 15, asgari: int = 3, sapma: Decimal = Decimal(20)) -> dict:
    dosyalar, uy = oku(yol)
    r = analiz(dosyalar, bugun, sla, asgari, sapma)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, bugun, sla)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Eksperleri rapor süresi, tespit–ödeme farkı ve itiraz oranına göre karşılaştırır.")
    p.add_argument("--dosyalar", type=Path, help="Dosya No, Eksper, Branş, Atama Tarihi, Rapor Tarihi, Tespit Tutarı, Ödenen Tutar, İtiraz, Eksper Ücreti")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--sla", type=int, default=15, help="Hedef rapor süresi, gün (varsayılan 15)")
    p.add_argument("--min", type=int, default=3, dest="asgari", help="Değerlendirme için en az raporlanan dosya (varsayılan 3)")
    p.add_argument("--sapma", default="20", help="Genel değerden yüzde kaç fazlası işaretlensin (varsayılan 20)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "eksper_performans_analizi.xlsx")
    a = p.parse_args(argv)
    ornek = a.dosyalar is None
    yol = ORNEK / "ekspertiz_dosyalari.csv" if ornek else a.dosyalar
    bugun, sapma = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today(), para(a.sapma)
    if bugun is None or a.sla < 1 or a.asgari < 1 or sapma is None or sapma < 0:
        print("[X] --tarih GG.AA.YYYY; --sla ve --min en az 1; --sapma eksi olamaz")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, bugun, a.sla, a.asgari, sapma)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    g = r["genel"]
    print(f"[OK] {len(r['eksperler'])} eksper · {g['dosya']} dosya · raporlanan {g['raporlanan']} · bekleyen {g['bekleyen']}" + (f" · ortalama süre {g['ort_sure']} gün" if g["ort_sure"] is not None else "")
          + (f" · mutlak sapma %{g['mutlak']}" if g["mutlak"] is not None else "") + (f" · itiraz %{g['itiraz_oran']}" if g["itiraz_oran"] is not None else ""))
    for e in r["eksperler"]:
        if e["durum"] not in ("Olağan", "Az veri"):
            print(f"[!] {e['eksper']}: {e['durum']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
