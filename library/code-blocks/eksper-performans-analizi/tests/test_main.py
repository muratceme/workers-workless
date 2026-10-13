import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)


def d(no, eksper, sure, tespit=None, odenen=None, itiraz=False, ucret=None, atama=40):
    """sure: rapor kaç günde geldi (None = rapor yok); atama: kaç gün önce."""
    at = BUGUN - timedelta(days=atama)
    return {"no": no, "eksper": eksper, "brans": "", "atama": at, "rapor": None if sure is None else at + timedelta(days=sure), "tespit": None if tespit is None else D(tespit), "odenen": None if odenen is None else D(odenen),
            "itiraz": itiraz, "ucret": None if ucret is None else D(ucret)}


def veri():
    return [d("1", "A", 10, 100, 100, ucret=10), d("2", "A", 20, 200, 180, True, 20), d("3", "a", 30, 300), d("4", "B", 5, 1000, 1100), d("5", "B", None, atama=20)]


class AnalizTesti(unittest.TestCase):
    def test_gostergeler(self):
        r = main.analiz(veri(), BUGUN, asgari=1)
        k = lambda e: (e["dosya"], e["raporlanan"], e["ort_sure"], e["medyan"], e["asim"], e["asim_oran"], e["bekleyen"], e["enb_bekleme"], e["tespit"], e["odenen"], e["fark"], e["mutlak"], e["itiraz"], e["itiraz_oran"], e["ucret"], e["ort_ucret"])   # noqa: E731
        self.assertEqual([(e["eksper"], e["durum"]) + k(e) for e in r["eksperler"]],
                         [("A", "Süre uzun", 3, 3, D("20.00"), D(20), 2, D("66.67"), 0, None, D(300), D(280), D("-6.67"), D("6.67"), 1, D("33.33"), D(30), D("15.00")),
                          ("B", "Olağan", 2, 1, D("5.00"), D(5), 0, D("0.00"), 1, 20, D(1000), D(1100), D("10.00"), D("10.00"), 0, D("0.00"), D(0), None)])
        self.assertEqual(k(r["genel"]), (5, 4, D("16.25"), D("15.0"), 2, D("50.00"), 1, 20, D(1300), D(1380), D("6.15"), D("9.23"), 1, D("25.00"), D(30), D("15.00")))
        self.assertEqual(({x["no"]: (x["sure"], x["bekleme"], x["fark"], x["fark_oran"]) for x in r["dosyalar"]}, [x["no"] for x in r["bekleyenler"]]),
                         ({"1": (10, None, D(0), D("0.00")), "2": (20, None, D(-20), D("-10.00")), "3": (30, None, None, None), "4": (5, None, D(100), D("10.00")), "5": (None, 20, None, None)}, ["5"]))
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]), [("Eksper genel değerlerden sapıyor", "A"), ("Rapor gecikti", "5 · B")])

    def test_esikler(self):
        s = lambda **kw: {e["eksper"]: e["durum"] for e in main.analiz(veri(), BUGUN, **kw)["eksperler"]}   # noqa: E731
        self.assertEqual((s(), s(asgari=1, sapma=D(25)), s(asgari=1, sapma=D(5)), s(asgari=1, sla=20)["A"]),
                         ({"A": "Süre uzun", "B": "Az veri"}, {"A": "Olağan", "B": "Olağan"}, {"A": "Süre uzun", "B": "Tespit–ödeme farkı yüksek"}, "Süre uzun"))
        r = main.analiz(veri(), BUGUN, sla=20)
        self.assertEqual((r["genel"]["asim"], [u["tur"] for u in r["uyarilar"]]), (1, ["Eksper genel değerlerden sapıyor"]))
        l = [d(str(i), "X", 10, 100, 100, i < 2) for i in range(4)] + [d(f"y{i}", "Y", 10, 100, 100) for i in range(8)] + [d("z", "Z", 10, 100, 100, True), d("z2", "Z", 10, 100, 100), d("z3", "Z", 10, 100, 100)]
        r = main.analiz(l, BUGUN)
        self.assertEqual(({e["eksper"]: (e["durum"], e["itiraz_oran"]) for e in r["eksperler"]}, r["genel"]["itiraz_oran"]), ({"Y": ("Olağan", D("0.00")), "X": ("İtiraz yüksek", D("50.00")), "Z": ("Olağan", D("33.33"))}, D("20.00")))
        r = main.analiz([d("1", "A", -3, 100, 0)], BUGUN, asgari=1)
        self.assertEqual((r["dosyalar"][0]["sure"], r["dosyalar"][0]["fark"], [u["tur"] for u in r["uyarilar"]], r["eksperler"][0]["fark"]), (0, D(-100), ["Rapor tarihi atamadan önce"], D("-100.00")))
        with self.assertRaises(ValueError):
            main.analiz([], BUGUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "ekspertiz_dosyalari.csv", cikti, BUGUN)
            self.assertEqual([(e["eksper"], e["durum"], e["dosya"], e["raporlanan"], e["ort_sure"], e["asim"], e["bekleyen"], e["tespit"], e["odenen"], e["fark"], e["mutlak"], e["itiraz_oran"], e["ucret"]) for e in r["eksperler"]],
                             [("Eksper A", "Olağan", 6, 6, D("9.00"), 0, 0, D(495000), D(483000), D("-2.42"), D("2.42"), D("0.00"), D(21500)),
                              ("Eksper B", "Süre uzun; Tespit–ödeme farkı yüksek; İtiraz yüksek", 5, 4, D("21.00"), 4, 1, D(278000), D(239000), D("-14.03"), D("14.03"), D("50.00"), D(11500)),
                              ("Eksper C", "Az veri", 2, 2, D("10.00"), 0, 0, D(172000), D(172000), D("0.00"), D("0.00"), D("0.00"), D(8000))])
            g = r["genel"]
            self.assertEqual((g["dosya"], g["raporlanan"], g["ort_sure"], g["asim_oran"], g["tespit"], g["odenen"], g["fark"], g["mutlak"], g["itiraz_oran"], g["enb_bekleme"]),
                             (13, 12, D("13.17"), D("33.33"), D(945000), D(894000), D("-5.40"), D("5.40"), D("16.67"), 24))
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]], [("Orta", "Eksper genel değerlerden sapıyor", "Eksper B"), ("Orta", "Rapor gecikti", "E-11 · Eksper B")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Eksperler", "Bekleyen Raporlar", "Dosyalar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "e.csv"
            y.write_text("Hasar No;Eksper Adı;Görevlendirme Tarihi;Rapor Teslim Tarihi;Ekspertiz Tutarı;Ödeme Tutarı;İtiraz Var mı;Ücret\nD1;Ali;01.09.2026;10.09.2026;1.500,50;0;EVET;250\n;Veli;02.09.2026;;;;;\nD3;Ali;x;;;;;\n"
                         "D4;Ali;03.09.2026;y;;;;\nD5;Ali;03.09.2026;;-1;;;\nD6;;03.09.2026;;;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["eksper"], x["atama"], x["rapor"], x["tespit"], x["odenen"], x["itiraz"], x["ucret"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("D1", "Ali", date(2026, 9, 1), date(2026, 9, 10), D("1500.50"), D(0), True, D(250)), ("Satır 3", "Veli", date(2026, 9, 2), None, None, None, False, None)],
                              [("Dosya okunamadı", "D3"), ("Dosya okunamadı", "D4"), ("Dosya okunamadı", "D5")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--dosyalar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "12.10.2026", "--sla", "10", "--min", "1", "--sapma", "30"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--sla", "0"], ["--min", "0"], ["--sapma", "-1"])], [2] * 4)
            self.assertEqual(main.main(["--dosyalar", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Dosya No;Eksper\nA;B\n", "Eksper;Atama Tarihi\n;01.03.2026\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
