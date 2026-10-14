import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def g(gun, bolum, pol, yatis, ameliyat, dolu, yatak, gelir, h_pol=None, h_gelir=None):
    c = lambda x: None if x is None else D(x)   # noqa: E731
    return {"tarih": date(2026, 10, gun), "bolum": bolum, "pol": D(pol), "yatis": D(yatis), "ameliyat": D(ameliyat), "dolu": c(dolu), "yatak": c(yatak), "gelir": D(gelir), "acil": None,
            "h_pol": c(h_pol), "h_gelir": c(h_gelir)}


def veri():
    return [g(5, "B1", 100, 10, 5, 18, 20, 50000, 120, 60000), g(6, "b1", 80, 6, 3, 20, 20, 40000, 120, 60000), g(5, "B2", 50, 2, 0, 4, 10, 10000)]


def k(x):
    return (x["pol"], x["yatis"], x["ameliyat"], x["dolu"], x["yatak"], x["doluluk"], x["yatis_oran"], x["hasta_gelir"], x["pol_hedef"], x["gelir_hedef"])


class PanoTesti(unittest.TestCase):
    def test_gostergeler(self):
        r = main.pano(veri())
        self.assertEqual([(b["bolum"], b["gelir_pay"]) + k(b) for b in r["bolumler"]],
                         [("B1", D("90.00"), D(180), D(16), D(8), D(38), D(40), D("95.00"), D("8.89"), D("459.18"), D("75.00"), D("75.00")),
                          ("B2", D("10.00"), D(50), D(2), D(0), D(4), D(10), D("40.00"), D("4.00"), D("192.31"), None, None)])
        self.assertEqual(k(r["toplam"]), (D(230), D(18), D(8), D(42), D(50), D("84.00"), D("7.83"), D("403.23"), D("75.00"), D("75.00")))
        self.assertEqual([(x["tarih"].day, x["n"], x["pol"]) for x in r["gunler"]], [(5, 2, D(150)), (6, 1, D(80))])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Bilgi", "Doluluk düşük", "B2"), ("Orta", "Gelir hedefinin gerisinde", "B1"), ("Orta", "Poliklinik hedefinin gerisinde", "B1")])

    def test_esikler(self):
        r = main.pano(veri(), yuksek=D(90), dusuk=D(30), hedef=D(70))
        self.assertEqual([(u["tur"], u["kim"]) for u in r["uyarilar"]], [("Doluluk çok yüksek", "B1")])
        self.assertIsNone(main.pano([g(1, "X", 0, 0, 0, None, None, 0)])["toplam"]["doluluk"])
        with self.assertRaises(ValueError):
            main.pano([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "bolum_gunluk_verileri.csv", cikti)
            self.assertEqual([(b["bolum"], b["doluluk"], b["pol_hedef"], b["gelir_hedef"]) for b in r["bolumler"]],
                             [("Genel Cerrahi", D("82.14"), D("99.09"), D("100.15")), ("Kardiyoloji", D("74.11"), D("86.21"), D("82.69")), ("Ortopedi", D("69.64"), D("98.00"), D("97.20")),
                              ("Kadın Doğum", D("44.84"), D("95.56"), D("95.39")), ("Dahiliye", D("95.24"), D("98.07"), D("100.04")), ("Çocuk Sağlığı", D("67.35"), D("99.64"), D("101.15"))])
            self.assertEqual(k(r["toplam"]), (D(5522), D(419), D(167), D(1282), D(1708), D("75.06"), D("7.59"), D("2308.79"), D("96.01"), D("94.04")))
            self.assertEqual([(u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Dolu yatak yatak sayısından fazla", "Dahiliye · 03.10.2026"), ("Doluluk çok yüksek", "Dahiliye"), ("Gelir hedefinin gerisinde", "Kardiyoloji"),
                              ("Poliklinik hedefinin gerisinde", "Kardiyoloji"), ("Doluluk düşük", "Kadın Doğum")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Bölümler", "Günler", "Günlük Pano", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "v.csv"
            y.write_text("Gün;Branş;Muayene;Yatış Sayısı;Dolu Yatak;Yatak Sayısı;Ciro\n01.10.2026;A;1.200;10;22;20;1.500.000\n01.10.2026;a;1;;;;\nx;A;1;;;;\n02.10.2026;A;-1;;;;\n02.10.2026;;1;;;;\n",
                         encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["bolum"], x["pol"], x["gelir"], x["dolu"]) for x in l], sorted(u["tur"] for u in uy)),
                             ([("A", D(1200), D(1500000), D(22))], ["Dolu yatak yatak sayısından fazla", "Gün iki kez yazılmış", "Satır okunamadı", "Satır okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--veriler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--yuksek", "90", "--dusuk", "60", "--hedef", "95"]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--yuksek", "x"], ["--dusuk", "96"], ["--hedef", "-1"])], [2] * 3)
            self.assertEqual(main.main(["--veriler", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Gün;Muayene\n01.10.2026;5\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
