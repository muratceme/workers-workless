"""
Hastane Operasyon Panosu — Workers / Workless kod bloğu
Hastane Yönetimi › Hastane Müdürü
Poliklinik, yatış, ameliyat, doluluk ve gelir göstergelerini bölüm ve gün bazında panoda toplar.
  Yatak doluluk %       = Σ dolu yatak (gün sonu sayımı) ÷ Σ yatak sayısı × 100   (yatak sayısı yazılmış günler)
  Poliklinikten yatış % = yatış ÷ poliklinik muayenesi × 100
  Hasta başına gelir    = gelir ÷ (poliklinik + yatış)
  Hedef gerçekleşme %   = gerçekleşen ÷ hedef × 100   (poliklinik ve gelir hedefi yazılmış günler üzerinden)
  Uyarı                 = doluluk --yuksek (%95) üstü ya da --dusuk (%50) altı · hedef gerçekleşmesi --hedef (%90) altı · dolu yatak yatak sayısından fazla
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"tarih": ("tarih", "gun"), "bolum": ("bolum", "brans", "klinik", "servis"), "pol": ("poliklinik", "poliklinik muayene", "muayene", "muayene sayisi"),
         "yatis": ("yatis", "yatis sayisi", "yatan hasta"), "ameliyat": ("ameliyat", "ameliyat sayisi", "operasyon"), "dolu": ("dolu yatak", "dolu yatak sayisi", "yatan"),
         "yatak": ("yatak", "yatak sayisi", "aktif yatak"), "gelir": ("gelir", "ciro", "faturalanan", "hasilat"), "acil": ("acil", "acil basvuru", "acil muayene"),
         "h_pol": ("hedef poliklinik", "poliklinik hedefi"), "h_gelir": ("hedef gelir", "gelir hedefi", "butce gelir")}
SAYI = ("pol", "yatis", "ameliyat", "dolu", "yatak", "gelir", "acil", "h_pol", "h_gelir")
SIFIR = Decimal(0)
GUN_AD = ("Pzt", "Sal", "Çar", "Per", "Cum", "Cmt", "Paz")


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("tarih", "bolum")):
        bolum = metin(r["bolum"])
        if not bolum:
            continue
        t, d = tarih(r["tarih"]), {a: para(r[a]) if metin(r[a]) else None for a in SAYI}
        if t is None or any(metin(r[a]) and (d[a] is None or d[a] < 0) for a in SAYI):
            uy.append(uyari("Yüksek", "Satır okunamadı", f"{bolum} · {metin(r['tarih'])}", f"Satır {r['_satir']}: tarih ya da sayılardan biri okunamadı. Satır alınmadı."))
            continue
        if (t, katla(bolum)) in gorulen:
            uy.append(uyari("Orta", "Gün iki kez yazılmış", f"{bolum} · {t:%d.%m.%Y}", f"Satır {r['_satir']}: ikinci satır alınmadı."))
            continue
        gorulen.add((t, katla(bolum)))
        if d["dolu"] is not None and d["yatak"] is not None and d["dolu"] > d["yatak"]:
            uy.append(uyari("Orta", "Dolu yatak yatak sayısından fazla", f"{bolum} · {t:%d.%m.%Y}", f"Dolu {d['dolu']:,.0f}, yatak {d['yatak']:,.0f}; ek yatak ya da kayıt hatası olabilir."))
        l.append({"tarih": t, "bolum": bolum, **d})
    return l, uy


def _oran(pay, payda, kat=100) -> Decimal | None:
    return yuvarla(Decimal(pay) * kat / Decimal(payda)) if payda else None


def gosterge(l: list[dict]) -> dict:
    def top(a, kosul=None):
        return sum((g[a] for g in l if g[a] is not None and (kosul is None or g[kosul] is not None)), SIFIR)
    pol, yatis, gelir = top("pol"), top("yatis"), top("gelir")
    return {"gun": len({g["tarih"] for g in l}), "pol": pol, "yatis": yatis, "ameliyat": top("ameliyat"), "acil": top("acil"), "gelir": gelir, "dolu": top("dolu", "yatak"), "yatak": top("yatak", "dolu"),
            "doluluk": _oran(top("dolu", "yatak"), top("yatak", "dolu")), "yatis_oran": _oran(yatis, pol), "hasta_gelir": _oran(gelir, pol + yatis, 1),
            "h_pol": top("h_pol"), "pol_hedef": _oran(top("pol", "h_pol"), top("h_pol")), "h_gelir": top("h_gelir"), "gelir_hedef": _oran(top("gelir", "h_gelir"), top("h_gelir"))}


def pano(gunler: list[dict], yuksek: Decimal = Decimal(95), dusuk: Decimal = Decimal(50), hedef: Decimal = Decimal(90)) -> dict:
    if not gunler:
        raise ValueError("Panoya girecek kayıt yok.")
    uy, bg, tg = [], {}, {}
    for g in gunler:
        bg.setdefault(katla(g["bolum"]), []).append(g)
        tg.setdefault(g["tarih"], []).append(g)
    toplam = gosterge(gunler)
    bolumler = sorted(({"bolum": l[0]["bolum"], **gosterge(l)} for l in bg.values()), key=lambda b: (-b["gelir"], katla(b["bolum"])))
    for b in bolumler:
        b["gelir_pay"] = _oran(b["gelir"], toplam["gelir"])
        if b["doluluk"] is not None and b["doluluk"] > yuksek:
            uy.append(uyari("Orta", "Doluluk çok yüksek", b["bolum"], f"Ortalama doluluk %{b['doluluk']} (eşik %{yuksek}); yatış bekleyen hasta ve acil birikmesi riski."))
        elif b["doluluk"] is not None and b["doluluk"] < dusuk:
            uy.append(uyari("Bilgi", "Doluluk düşük", b["bolum"], f"Ortalama doluluk %{b['doluluk']} (eşik %{dusuk})."))
        for alan, ad in (("pol_hedef", "Poliklinik"), ("gelir_hedef", "Gelir")):
            if b[alan] is not None and b[alan] < hedef:
                uy.append(uyari("Orta", f"{ad} hedefinin gerisinde", b["bolum"], f"Hedef gerçekleşmesi %{b[alan]} (eşik %{hedef})."))
    detay = sorted(({"tarih": g["tarih"], "bolum": g["bolum"], **gosterge([g])} for g in gunler), key=lambda g: (g["tarih"], katla(g["bolum"])))
    return {"bolumler": bolumler, "gunler": [{"tarih": t, "n": len(l), **gosterge(l)} for t, l in sorted(tg.items())], "detay": detay, "toplam": toplam, "uyarilar": uy}


BASLIK = ["Poliklinik", "Acil", "Yatış", "Poliklinikten Yatış %", "Ameliyat", "Dolu Yatak (gün)", "Yatak (gün)", "Doluluk %", "Gelir", "Hasta Başına Gelir", "Poliklinik Hedef %", "Gelir Hedef %"]


def _deger(g: dict) -> list:
    return [g["pol"], g["acil"] or None, g["yatis"], g["yatis_oran"], g["ameliyat"], g["dolu"] or None, g["yatak"] or None, g["doluluk"], g["gelir"], g["hasta_gelir"], g["pol_hedef"], g["gelir_hedef"]]


def rapor_yaz(cikti: Path, r: dict, yuksek: Decimal, dusuk: Decimal, hedef: Decimal) -> None:
    wb = Workbook()

    def satir(ws, on, g, kalin=False):
        n = len(on)
        k = _satir(ws, on + _deger(g), para_s=[n + i for i in (4, 8, 9, 10, 11, 12)],
                   tarih_s=(1,) if isinstance(on[0], date) else (),
                   dolgu={n + 8: None if g["doluluk"] is None else KIRMIZI if g["doluluk"] > yuksek else SARI if g["doluluk"] < dusuk else YESIL,
                          n + 11: None if g["pol_hedef"] is None else KIRMIZI if g["pol_hedef"] < hedef else YESIL if g["pol_hedef"] >= 100 else None,
                          n + 12: None if g["gelir_hedef"] is None else KIRMIZI if g["gelir_hedef"] < hedef else YESIL if g["gelir_hedef"] >= 100 else None})
        if kalin:
            for h in ws[k]:
                h.font = Font(bold=True)
    ws = wb.active
    ws.title = "Bölümler"
    _baslik(ws, ["Bölüm", "Gün", "Gelir Payı %"] + BASLIK, (22, 6, 9, 10, 8, 8, 11, 9, 11, 10, 9, 15, 12, 11, 11), "B2")
    for b in r["bolumler"]:
        satir(ws, [b["bolum"], b["gun"], b["gelir_pay"]], b)
    satir(ws, ["HASTANE", r["toplam"]["gun"], None], r["toplam"], True)
    ws.append([])
    ws.append([f"Doluluk = dolu yatak ÷ yatak sayısı (gün sonu sayımı). Eşikler: doluluk %{dusuk}–%{yuksek}, hedef %{hedef}."])
    ws = wb.create_sheet("Günler")
    _baslik(ws, ["Tarih", "Gün", "Bölüm Sayısı"] + BASLIK, (11, 6, 7, 10, 8, 8, 11, 9, 11, 10, 9, 15, 12, 11, 11), "B2")
    for g in r["gunler"]:
        satir(ws, [g["tarih"], GUN_AD[g["tarih"].weekday()], g["n"]], g)
    ws = wb.create_sheet("Günlük Pano")
    _baslik(ws, ["Tarih", "Gün", "Bölüm"] + BASLIK, (11, 6, 20, 10, 8, 8, 11, 9, 11, 10, 9, 15, 12, 11, 11), "D2")
    for g in r["detay"]:
        satir(ws, [g["tarih"], GUN_AD[g["tarih"].weekday()], g["bolum"]], g)
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, yuksek: Decimal = Decimal(95), dusuk: Decimal = Decimal(50), hedef: Decimal = Decimal(90)) -> dict:
    gunler, uy = oku(yol)
    r = pano(gunler, yuksek, dusuk, hedef)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, yuksek, dusuk, hedef)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Poliklinik, yatış, ameliyat, doluluk ve gelir göstergelerini bölüm bazında günlük panoda toplar.")
    p.add_argument("--veriler", type=Path, help="Tarih, Bölüm, Poliklinik, Acil, Yatış, Ameliyat, Dolu Yatak, Yatak Sayısı, Gelir, Hedef Poliklinik, Hedef Gelir")
    p.add_argument("--yuksek", default="95", help="Doluluk bu yüzdenin üstündeyse uyarılır (varsayılan 95)")
    p.add_argument("--dusuk", default="50", help="Doluluk bu yüzdenin altındaysa bilgi verilir (varsayılan 50)")
    p.add_argument("--hedef", default="90", help="Hedef gerçekleşmesi bu yüzdenin altındaysa uyarılır (varsayılan 90)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "hastane_operasyon_panosu.xlsx")
    a = p.parse_args(argv)
    yol = a.veriler or ORNEK / "bolum_gunluk_verileri.csv"
    yuksek, dusuk, hedef = para(a.yuksek), para(a.dusuk), para(a.hedef)
    if None in (yuksek, dusuk, hedef) or not 0 <= dusuk <= yuksek or hedef < 0:
        print("[X] 0 ≤ --dusuk ≤ --yuksek olmalı; --hedef sıfır ya da üstü")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, yuksek, dusuk, hedef)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {len(r['bolumler'])} bölüm · {t['gun']} gün · poliklinik {t['pol']:,.0f} · yatış {t['yatis']:,.0f} · ameliyat {t['ameliyat']:,.0f}"
          + (f" · doluluk %{t['doluluk']}" if t["doluluk"] is not None else "") + f" · gelir {t['gelir']:,.2f}")
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
