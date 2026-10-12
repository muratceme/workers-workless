"""
Proje Maliyet Kontrol Raporu — Workers / Workless kod bloğu
Teknik Ofis › Teknik Ofis Şefi

Bütçe, gerçekleşen maliyet ve hakediş gelirlerinden iş kalemi bazında maliyet sapması ve tamamlanma tahmini üretir.
  Kazanılan değer          = bütçe × fiziksel tamamlanma %
  Maliyet sapması          = kazanılan değer − gerçekleşen maliyet               (eksi: yapılan iş için bütçeden fazla harcanmış)
  Maliyet performansı      = kazanılan değer ÷ gerçekleşen maliyet
  Tamamlanma tahmini       = --yontem performans (varsayılan): gerçekleşen maliyet × bütçe ÷ kazanılan değer   (bugünkü verim sürerse)
                             --yontem butce: gerçekleşen maliyet + (bütçe − kazanılan değer)                   (kalan iş bütçesinde biterse)
                             tamamlanmış kalemde gerçekleşen maliyet; ilerlemesi veya maliyeti olmayan kalemde bütçe yöntemi
  Tamamlanma sapması       = bütçe − tamamlanma tahmini                          (eksi: bütçe aşımı bekleniyor)
  Bugüne kadarki brüt kâr  = hakediş geliri − gerçekleşen maliyet
  Durum                    = Başlamadı · Bütçe aşıldı (maliyet > bütçe) · Aşım bekleniyor (sapma < −--tolerans %5) · Bütçe altında (sapma > +tolerans) · Bütçede
Tahmin, girdiğiniz tamamlanma yüzdesine dayanır; yüzde iyimserse tahmin de iyimser olur. Taahhüt edilmiş ama faturalanmamış maliyetleri gerçekleşene eklemeniz gerekir.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"kalem": ("is kalemi", "kalem", "is grubu", "maliyet kalemi", "aciklama"), "butce": ("butce", "butce maliyeti", "planlanan maliyet"), "maliyet": ("gerceklesen maliyet", "gerceklesen", "fiili maliyet", "maliyet"),
         "oran": ("tamamlanma", "tamamlanma %", "tamamlanma yuzdesi", "fiziksel ilerleme", "ilerleme"), "hakedis": ("hakedis geliri", "hakedis", "gelir", "kesilen hakedis")}
SIRA = ("Bütçe aşıldı", "Aşım bekleniyor", "Bütçede", "Bütçe altında", "Başlamadı")
Z = Decimal(0)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("kalem", "butce", "maliyet", "oran")):
        k = metin(r["kalem"])
        if not k:
            continue
        kim = f"{k} (satır {r['_satir']})"
        b, m, o, h = (para(r[a]) for a in ("butce", "maliyet", "oran", "hakedis"))
        if b is None or b <= 0:
            uy.append(uyari("Yüksek", "Bütçe okunamadı", kim, f"'{metin(r['butce'])}'. Kalem alınmadı."))
            continue
        if m is None and metin(r["maliyet"]) or (o is None and metin(r["oran"])) or (m or Z) < 0 or not 0 <= (o or Z) <= 100:
            uy.append(uyari("Yüksek", "Maliyet veya tamamlanma okunamadı", kim, f"Gerçekleşen '{metin(r['maliyet'])}', tamamlanma '{metin(r['oran'])}'. Tamamlanma 0-100 arasında olmalı; kalem alınmadı."))
            continue
        if h is None and metin(r["hakedis"]):
            uy.append(uyari("Orta", "Hakediş okunamadı", kim, f"'{metin(r['hakedis'])}'. Boş sayıldı."))
        l.append({"kalem": k, "butce": b, "maliyet": m or Z, "oran": o or Z, "hakedis": h})
    return l, uy


def degerlendir(kalemler: list[dict], yontem: str = "performans", tolerans: Decimal = Decimal(5), sozlesme: Decimal | None = None) -> dict:
    if not kalemler:
        raise ValueError("Değerlendirilecek iş kalemi yok.")
    uy = []
    for k in kalemler:
        b, ac = k["butce"], k["maliyet"]
        ev = yuvarla(b * k["oran"] / 100)
        if k["oran"] >= 100:
            eac = ac
        elif yontem == "performans" and ev > 0 and ac > 0:
            eac = yuvarla(ac * b / ev)
        else:
            eac = ac + (b - ev)
        k |= {"ev": ev, "cv": ev - ac, "cpi": yuvarla(ev / ac) if ac else None, "eac": eac, "etc": eac - ac, "vac": b - eac, "vac_yuzde": yuvarla((b - eac) / b * 100),
              "kar": k["hakedis"] - ac if k["hakedis"] is not None else None}
        k["durum"] = ("Başlamadı" if not ac and not k["oran"] else "Bütçe aşıldı" if ac > b else "Aşım bekleniyor" if k["vac_yuzde"] < -tolerans else "Bütçe altında" if k["vac_yuzde"] > tolerans else "Bütçede")
        if ac > 0 and k["oran"] == 0:
            uy.append(uyari("Orta", "Maliyet var, ilerleme yok", k["kalem"], f"Gerçekleşen maliyet {ac:,.2f} ama tamamlanma %0. Tahmin bütçe yöntemiyle yapıldı; ilerleme yüzdesini kontrol edin."))
        if k["oran"] >= 100 and k["hakedis"] is not None and k["hakedis"] < ac:
            uy.append(uyari("Yüksek", "Tamamlanan kalem zararda", k["kalem"], f"Hakediş {k['hakedis']:,.2f}, maliyet {ac:,.2f}."))
    t = {a: sum((k[a] for k in kalemler), Z) for a in ("butce", "maliyet", "ev", "cv", "eac", "etc", "vac")}
    t["hakedis"] = sum((k["hakedis"] for k in kalemler if k["hakedis"] is not None), Z)
    t |= {"oran": yuvarla(t["ev"] / t["butce"] * 100), "cpi": yuvarla(t["ev"] / t["maliyet"]) if t["maliyet"] else None, "vac_yuzde": yuvarla(t["vac"] / t["butce"] * 100), "kar": t["hakedis"] - t["maliyet"],
          "sozlesme": sozlesme, "tahmini_kar": sozlesme - t["eac"] if sozlesme is not None else None, "butce_kar": sozlesme - t["butce"] if sozlesme is not None else None}
    t["tahmini_marj"] = yuvarla(t["tahmini_kar"] / sozlesme * 100) if sozlesme else None
    return {"kalemler": sorted(kalemler, key=lambda k: (SIRA.index(k["durum"]), k["vac"])), "toplam": t, "yontem": yontem, "tolerans": tolerans, "uyarilar": uy,
            "say": {s: sum(k["durum"] == s for k in kalemler) for s in SIRA}}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    renk = {"Bütçe aşıldı": KIRMIZI, "Aşım bekleniyor": KIRMIZI, "Bütçe altında": YESIL, "Başlamadı": MAVI}
    t = r["toplam"]
    ws = wb.active
    ws.title = "Maliyet Kontrol"
    _baslik(ws, ["Durum", "İş Kalemi", "Bütçe", "Tamamlanma %", "Kazanılan Değer", "Gerçekleşen Maliyet", "Maliyet Sapması", "Maliyet Performansı", "Tamamlanma Tahmini", "Kalan Maliyet", "Tamamlanma Sapması",
                 "Sapma %", "Hakediş Geliri", "Bugüne Kadarki Brüt Kâr", "Açıklama / Aksiyon (siz doldurun)"], (17, 30, 16, 10, 16, 16, 15, 10, 16, 16, 16, 9, 16, 16, 34))
    for k in [*r["kalemler"], {"durum": "", "kalem": "TOPLAM"} | t]:
        n = _satir(ws, [k["durum"], k["kalem"], k["butce"], k["oran"], k["ev"], k["maliyet"], k["cv"], k["cpi"], k["eac"], k["etc"], k["vac"], k["vac_yuzde"], k["hakedis"], k["kar"], ""],
                   para_s=tuple(range(3, 15)), dolgu={1: renk.get(k["durum"]), 7: KIRMIZI if k["cv"] < 0 else None, 11: KIRMIZI if k["vac"] < 0 else None})
    for c in ws[n]:
        c.font = Font(bold=True)
    ws.append([])
    if t["sozlesme"] is not None:
        ws.append([f"Sözleşme bedeli {t['sozlesme']:,.2f} · bütçeye göre kâr {t['butce_kar']:,.2f} · tamamlanma tahminine göre kâr {t['tahmini_kar']:,.2f} (%{t['tahmini_marj']})."])
    ws.append([("Tamamlanma tahmini: gerçekleşen maliyet × bütçe ÷ kazanılan değer (bugünkü verim sürerse)." if r["yontem"] == "performans" else "Tamamlanma tahmini: gerçekleşen maliyet + kalan işin bütçesi.")
               + f" Durum eşiği %{r['tolerans'].normalize():f}. Tahmin, girilen fiziksel tamamlanma yüzdesine dayanır."])
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, **ayar) -> dict:
    kalemler, uy = oku(yol)
    r = degerlendir(kalemler, **ayar)
    uy += r["uyarilar"]
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="İş kalemi bazında maliyet sapması ve tamamlanma tahmini üretir.")
    p.add_argument("--kalemler", type=Path, default=ORNEK / "is_kalemleri.csv", help="İş Kalemi, Bütçe, Gerçekleşen Maliyet, Tamamlanma %%, Hakediş Geliri")
    p.add_argument("--yontem", choices=("performans", "butce"), default="performans", help="Tamamlanma tahmini yöntemi (varsayılan performans)")
    p.add_argument("--tolerans", default="5", help="Bütçeden bu %% kadar sapma 'Bütçede' sayılır (varsayılan 5)")
    p.add_argument("--sozlesme-bedeli", help="İsteğe bağlı: toplam sözleşme bedeli (tahmini proje kârı için; örnek veride 28.000.000)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "maliyet_kontrol_raporu.xlsx")
    a = p.parse_args(argv)
    sb = a.sozlesme_bedeli or ("28000000" if a.kalemler == ORNEK / "is_kalemleri.csv" else None)
    t, s = para(a.tolerans), para(sb) if sb else None
    if t is None or t < 0 or (sb and (s is None or s <= 0)):
        print("[X] --tolerans sıfır veya pozitif; --sozlesme-bedeli pozitif sayı olmalı")
        return 2
    if not a.kalemler.exists():
        print(f"[X] Dosya bulunamadı: {a.kalemler}")
        return 1
    try:
        r = calistir(a.kalemler, a.cikti, yontem=a.yontem, tolerans=t, sozlesme=s)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    x = r["toplam"]
    print(f"[OK] {len(r['kalemler'])} kalem · tamamlanma %{x['oran']} · bütçe {x['butce']:,.2f} · gerçekleşen {x['maliyet']:,.2f} · maliyet sapması {x['cv']:+,.2f}")
    print(f"[{'!' if x['vac'] < 0 else 'OK'}] Tamamlanma tahmini {x['eac']:,.2f} · bütçeye göre {x['vac']:+,.2f} (%{x['vac_yuzde']})" + (f" · tahmini kâr {x['tahmini_kar']:,.2f}" if x["tahmini_kar"] is not None else ""))
    for k in r["kalemler"]:
        if k["durum"] in ("Bütçe aşıldı", "Aşım bekleniyor"):
            print(f"[!] {k['kalem']}: {k['durum']} · tahmin {k['eac']:,.2f}, bütçe {k['butce']:,.2f} ({k['vac']:+,.2f})")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
