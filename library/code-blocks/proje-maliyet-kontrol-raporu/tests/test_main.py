import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def kl(butce, maliyet, oran, hakedis=None, kalem="K"):
    return {"kalem": kalem, "butce": D(butce), "maliyet": D(maliyet), "oran": D(oran), "hakedis": None if hakedis is None else D(hakedis)}


def tek(*a, **kw):
    return main.degerlendir([kl(*a)], **kw)["kalemler"][0]


class DegerlendirmeTesti(unittest.TestCase):
    def test_kalem(self):
        k = tek(10000000, 6600000, 60, 7200000)
        self.assertEqual((k["ev"], k["cv"], k["cpi"], k["eac"], k["etc"], k["vac"], k["vac_yuzde"], k["kar"], k["durum"]),
                         (D(6000000), D(-600000), D("0.91"), D(11000000), D(4400000), D(-1000000), D("-10.00"), D(600000), "Aşım bekleniyor"))
        k = tek(10000000, 6600000, 60, yontem="butce")
        self.assertEqual((k["eac"], k["vac"], k["vac_yuzde"], k["durum"], k["kar"]), (D(10600000), D(-600000), D("-6.00"), "Aşım bekleniyor", None))
        k = tek(6000000, 1350000, 25)
        self.assertEqual((k["eac"], k["vac_yuzde"], k["durum"]), (D(5400000), D("10.00"), "Bütçe altında"))
        self.assertEqual([(k["eac"], k["durum"]) for k in (tek(3000000, 0, 0), tek(1000000, 1150000, 100), tek(1000, 500, 0), tek(1000, 0, 50), tek(1000, 1200, 80))],
                         [(D(3000000), "Başlamadı"), (D(1150000), "Bütçe aşıldı"), (D(1500), "Aşım bekleniyor"), (D(500), "Bütçe altında"), (D(1500), "Bütçe aşıldı")])
        self.assertEqual([tek(1000, m, 50)["durum"] for m in (525, 526, 475, 474)], ["Bütçede", "Aşım bekleniyor", "Bütçede", "Bütçe altında"])         # tam %5 bütçede
        self.assertEqual(tek(1000, 600, 50, tolerans=D(25))["durum"], "Bütçede")

    def test_toplam_ve_uyarilar(self):
        l = [kl(1000000, 1150000, 100, 1200000, "H"), kl(10000000, 6600000, 60, 7200000, "K"), kl(6000000, 1350000, 25, 1800000, "İ"), kl(4000000, 800000, 20, 900000, "M"), kl(3000000, 0, 0, 0, "E")]
        r = main.degerlendir(l, sozlesme=D(28000000))
        t = r["toplam"]
        self.assertEqual((t["butce"], t["maliyet"], t["ev"], t["cv"], t["oran"], t["cpi"], t["eac"], t["vac"], t["vac_yuzde"], t["hakedis"], t["kar"], t["butce_kar"], t["tahmini_kar"], t["tahmini_marj"]),
                         (D(24000000), D(9900000), D(9300000), D(-600000), D("38.75"), D("0.94"), D(24550000), D(-550000), D("-2.29"), D(11100000), D(1200000), D(4000000), D(3450000), D("12.32")))
        self.assertEqual(([k["kalem"] for k in r["kalemler"]], r["say"]), (["H", "K", "M", "İ", "E"], {"Bütçe aşıldı": 1, "Aşım bekleniyor": 1, "Bütçede": 1, "Bütçe altında": 1, "Başlamadı": 1}))
        self.assertEqual(main.degerlendir(l)["toplam"]["tahmini_kar"], None)
        r = main.degerlendir([kl(1000, 500, 0), kl(1000, 1100, 100, 1000)])
        self.assertEqual(sorted(u["tur"] for u in r["uyarilar"]), ["Maliyet var, ilerleme yok", "Tamamlanan kalem zararda"])
        with self.assertRaises(ValueError):
            main.degerlendir([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "p.xlsx"
            r = main.calistir(ORNEK / "is_kalemleri.csv", cikti, sozlesme=D(28000000))
            self.assertEqual(({k["kalem"].split()[0]: (k["durum"], k["eac"]) for k in r["kalemler"]}, r["toplam"]["eac"], r["toplam"]["tahmini_kar"], r["uyarilar"]),
                             ({"Hafriyat": ("Bütçe aşıldı", D(1150000)), "Kaba": ("Aşım bekleniyor", D(11000000)), "İnce": ("Bütçe altında", D(5400000)), "Mekanik": ("Bütçede", D(4000000)),
                               "Elektrik": ("Başlamadı", D(3000000))}, D(24550000), D(3450000), []))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Maliyet Kontrol", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "k.csv"
            y.write_text("Kalem;Bütçe;Gerçekleşen;Tamamlanma %;Hakediş\nA;1.000,00;400;%50;çok\nB;yok;1;1;\nC;1.000;100;150;\nD;1.000;;;\nE;1.000;-5;10;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual([(k["kalem"], k["butce"], k["maliyet"], k["oran"], k["hakedis"]) for k in l], [("A", D("1000.00"), D(400), D(50), None), ("D", D(1000), D(0), D(0), None)])
            self.assertEqual(sorted(u["tur"] for u in uy), ["Bütçe okunamadı", "Hakediş okunamadı", "Maliyet veya tamamlanma okunamadı", "Maliyet veya tamamlanma okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--kalemler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--yontem", "butce", "--tolerans", "10", "--sozlesme-bedeli", "3.000,00"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(arg + ["--sozlesme-bedeli", "0"]), 2)
            self.assertEqual(main.main(["--kalemler", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("İş Kalemi;Bütçe;Gerçekleşen Maliyet;Tamamlanma\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
