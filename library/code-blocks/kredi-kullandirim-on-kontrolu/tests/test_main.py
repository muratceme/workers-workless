import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 9)                  # Cuma


def onay(no="O1", limit=1000, risk=0, **kw):
    o = {"onay": no, "musteri_no": "M1", "limit": D(limit), "pb": "TL", "urunler": set(), "urunler_ad": "", "vade": None, "faiz": None, "bitis": date(2027, 1, 1), "risk": D(risk),
         "soz_tutar": D(5000), "soz_durum": "İmzalı", "soz_var": True}
    return o | kw


def talep(no="T1", tutar=100, onay_no="O1", **kw):
    t = {"no": no, "musteri_no": "M1", "musteri": "M", "sube": "", "onay": onay_no, "urun": "Rotatif", "tutar": D(tutar), "pb": "TL", "vade": D(12), "faiz": D(45), "tarih": BUGUN}
    return t | kw


def sart(metin="Şart", tamam=True, sonra=False, gereken=None, tesis=None, tarih=None):
    return {"sart": metin, "tur": "Teminat", "sonra": sonra, "zaman": "", "gereken": D(gereken) if gereken is not None else None, "tesis": D(tesis) if tesis is not None else None,
            "tamam": tamam, "durum": "Tamamlandı" if tamam else "Bekliyor", "tarih": tarih}


def bulgu(r, i=0):
    return [(b[0], b[2]) for b in r["talepler"][i]["bulgular"]]


def calis(t, o=None, s=None, **kw):
    return main.kontrol(t if isinstance(t, list) else [t], {"O1": o or onay()}, {"O1": s or []}, BUGUN, **kw)


class KontrolTesti(unittest.TestCase):
    def test_temiz_talep(self):
        r = calis(talep())
        self.assertEqual((r["talepler"][0]["sonuc"], r["talepler"][0]["bulgular"], r["talepler"][0]["kullanilabilir"]), ("Kullandırılabilir", [], D(1000)))

    def test_limit_ve_sirali_dusum(self):
        r = calis([talep("T1", 400), talep("T2", 300), talep("T3", 200), talep("T4", 100)], onay(risk=200))
        self.assertEqual([t["sonuc"] for t in r["talepler"]], ["Kullandırılabilir", "Kullandırılabilir", "Engel var", "Kullandırılabilir"])   # 800 − 400 − 300 = 100 kalır
        self.assertEqual([t["kullanilabilir"] for t in r["talepler"]], [D(800), D(400), D(100), D(100)])
        self.assertIn("aşım 100,00", r["talepler"][2]["bulgular"][0][3])
        self.assertEqual(bulgu(calis(talep(tutar=1000), onay(limit=1000))), [])                 # tam limit kullanılabilir

    def test_onay_kosullari(self):
        o = onay(urunler={"rotatif", "ticari taksitli"}, urunler_ad="Rotatif, Ticari Taksitli", vade=D(24), faiz=D(42), bitis=date(2026, 10, 20))
        self.assertEqual(bulgu(calis(talep(), o)), [("Kontrol", "Onayın geçerliliği dolmak üzere")])
        self.assertEqual(bulgu(calis(talep(), o, yaklasan=5)), [])
        r = calis(talep(urun="Spot", vade=D(25), faiz=D("41.9"), pb="USD", musteri_no="M2", tarih=date(2026, 10, 21)), o)
        self.assertEqual(sorted(x[1] for x in bulgu(r)), ["Faiz asgari oranın altında", "Müşteri onayla uyuşmuyor", "Onayın geçerliliği dolmuş", "Para birimi onaydan farklı",
                                                          "Vade azami vadeyi aşıyor", "Ürün onay kapsamında değil"])
        self.assertEqual(bulgu(calis(talep(vade=D(24), faiz=D(42), urun="TİCARİ TAKSİTLİ"), o, yaklasan=5)), [])   # sınır değerler geçerli

    def test_sozlesme(self):
        self.assertEqual(bulgu(calis(talep(), onay(soz_durum="İmza bekliyor"))), [("Engel", "Sözleşme imzalı değil")])
        self.assertEqual(bulgu(calis(talep(tutar=300), onay(risk=500, soz_tutar=D(700)))), [("Engel", "Sözleşme tutarı yetersiz")])
        self.assertEqual(bulgu(calis(talep(tutar=200), onay(risk=500, soz_tutar=D(700)))), [])
        self.assertEqual(bulgu(calis(talep(), onay(soz_var=False, soz_tutar=None, soz_durum=""))), [("Kontrol", "Sözleşme bilgisi girilmemiş")])

    def test_sartlar(self):
        s = [sart("İpotek", gereken=1000, tesis=600), sart("Teşvik belgesi", tamam=False), sart("Sigorta", tamam=False, sonra=True), sart("Kefalet"),
             sart("Rehin", tarih=date(2026, 10, 12))]
        r = calis([talep("T1"), talep("T2")], s=s)
        self.assertEqual(bulgu(r), [("Engel", "Teminat tutarı eksik"), ("Engel", "Kullandırım öncesi şart tamamlanmamış"), ("Kontrol", "Şartın tamamlanma tarihi kullandırımdan sonra"),
                                    ("Bilgi", "Kullandırım sonrası şart")])
        self.assertIn("eksik 400,00", r["talepler"][0]["bulgular"][0][3])
        self.assertEqual([(x["t"]["no"], x["s"]["sart"]) for x in r["takip"]], [("T1", "Sigorta")])     # aynı şart bir kez
        self.assertEqual(calis(talep(), s=[sart("Sigorta", tamam=False, sonra=True)])["talepler"][0]["sonuc"], "Kullandırılabilir")

    def test_tarih_ve_onaysiz(self):
        self.assertEqual(bulgu(calis(talep(tarih=date(2026, 10, 10)))), [("Kontrol", "Kullandırım tarihi hafta sonu")])
        self.assertEqual(bulgu(calis(talep(tarih=date(2026, 10, 7)))), [("Kontrol", "Kullandırım tarihi geçmiş")])
        r = calis(talep(onay_no="YOK"))
        self.assertEqual((bulgu(r), r["talepler"][0]["sonuc"]), ([("Engel", "Onay kaydı bulunamadı")], "Engel var"))


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "k.xlsx"
            r = main.calistir(ORNEK / "talepler.csv", ORNEK / "onaylar.csv", cikti, BUGUN, ORNEK / "sartlar.csv")
            d = {t["no"]: (t["sonuc"], sorted(b[2] for b in t["bulgular"] if b[0] == "Engel")) for t in r["talepler"]}
            self.assertEqual(d["KT-9001"], ("Kullandırılabilir", []))
            self.assertEqual(d["KT-9002"], ("Engel var", ["Limit yetersiz"]))                   # 20 − 12,5 − 5 = 2,5 milyon < 4 milyon
            self.assertEqual(d["KT-9003"], ("Engel var", ["Teminat tutarı eksik"]))
            self.assertEqual(d["KT-9005"][0], "Kontrol edin")
            self.assertEqual(d["KT-9006"][1], ["Kullandırım öncesi şart tamamlanmamış", "Sözleşme tutarı yetersiz"])
            self.assertEqual(d["KT-9008"][1], ["Sözleşme imzalı değil", "Vade azami vadeyi aşıyor"])
            self.assertEqual(sum(v[0] == "Engel var" for v in d.values()), 8)
            self.assertEqual(len(r["takip"]), 2)
            self.assertEqual(load_workbook(cikti).sheetnames, ["Talep Sonuçları", "Bulgular", "Şart Takibi", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            t, o, s = (Path(tmp) / x for x in ("t.csv", "o.csv", "s.csv"))
            t.write_text("Talep No;Onay No;Tutar;Kullandırım Tarihi\nT1;O1;1.000,00;09.10.2026\nT2;O1;sıfır;\nT3;O1;500;yarın\n", encoding="utf-8")
            o.write_text("Onay No;Onaylı Limit;Ürünler;Geçerlilik Bitiş;Sözleşme Durumu\nO1;10.000,00;*;bozuk;İmzalı\nO1;5;;;\nO2;yok;;;\n", encoding="utf-8")
            s.write_text("Onay No;Şart;Zamanlama;Durum\nO1;İpotek;Kullandırım öncesi;Tesis edildi\nO1;Sigorta;Sonra;\nO9;Kefalet;;\n", encoding="utf-8")
            tl_, on, sr, uy = main.oku(t, o, s)
            self.assertEqual(([x["no"] for x in tl_], on["O1"]["urunler"], on["O1"]["bitis"], on["O1"]["soz_var"]), (["T1", "T3"], set(), None, True))
            self.assertEqual([(x["sart"], x["tamam"], x["sonra"]) for x in sr["O1"]], [("İpotek", True, False), ("Sigorta", False, True)])
            self.assertEqual(sorted(u["tur"] for u in uy), ["Geçerlilik tarihi okunamadı", "Kullandırım tarihi okunamadı", "Onay satırı okunamadı", "Onay satırı okunamadı",
                                                            "Talep tutarı okunamadı", "Şartın onayı listede yok"])
            arg = ["--talepler", str(t), "--onaylar", str(o), "--sartlar", str(s), "--cikti", str(Path(tmp) / "x.xlsx"), "--tarih", "09.10.2026"]
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(arg[:-1] + ["dün"]), 2)
            self.assertEqual(main.main(["--onaylar", str(Path(tmp) / "yok.csv")]), 1)
            t.write_text("Talep No;Onay No;Tutar\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "o.xlsx")]), 0)


if __name__ == "__main__":
    unittest.main()
