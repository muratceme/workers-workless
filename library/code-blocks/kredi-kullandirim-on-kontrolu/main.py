"""
Kredi Kullandırım Ön Kontrolü — Workers / Workless kod bloğu
Bankacılık Operasyonları › Kredi Operasyon Uzmanı

Kullandırım taleplerini tahsis onayındaki şartlar, teminat tesis durumu ve sözleşme bilgileriyle karşılaştırır;
kullandırıma engel olan eksikleri listeler.
  Engel (kullandırım yapılmamalı):
    - Onay kaydı yok, müşteri numarası onayla uyuşmuyor, onayın geçerliliği dolmuş
    - Limit: mevcut risk + (aynı çalıştırmada aynı onaya bağlı önceki talepler) + talep > onaylı limit
    - Para birimi veya ürün onaydakinden farklı; vade azami vadeyi aşıyor; faiz asgari faizin altında
    - Sözleşme yok / imzasız; sözleşme tutarı (mevcut risk + talep) tutarını karşılamıyor
    - "Kullandırım öncesi" şart tamamlanmamış; teminat tesis edilen tutar istenen tutarın altında
  Kontrol edin: onayın geçerliliği --yaklasan gün içinde doluyor, kullandırım tarihi hafta sonu / geçmiş tarih
  Bilgi: "Kullandırım sonrası" şartlar (takip listesine alınır)
  Talepler dosyadaki sırayla işlenir; engeli olmayan her talep, aynı onaya bağlı sonraki talepler için limitten düşülür.
Kod yalnız girilen kayıtları karşılaştırır; teminatın hukuken geçerli tesis edildiğini veya belgelerin aslını doğrulamaz.
Nihai kullandırım yetkisi bankanızın onay akışındadır. İnternete bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: 10 talep
    python main.py --talepler talepler.xlsx --onaylar onaylar.xlsx --sartlar sartlar.xlsx --tarih 09.10.2026
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


TALEP_SUTUNLARI = {"no": ("talep no", "talep numarasi", "kullandirim no"), "musteri_no": ("musteri no", "musteri numarasi"), "musteri": ("musteri", "musteri adi", "unvan"),
                   "sube": ("sube",), "onay": ("onay no", "tahsis no", "karar no", "limit no"), "urun": ("urun", "kredi turu"), "tutar": ("tutar", "talep tutari", "kullandirim tutari"),
                   "pb": ("para birimi", "pb"), "vade": ("vade", "vade ay"), "faiz": ("faiz orani", "faiz", "faiz orani yuzde"), "tarih": ("kullandirim tarihi", "tarih", "valor")}
ONAY_SUTUNLARI = {"onay": ("onay no", "tahsis no", "karar no", "limit no"), "musteri_no": ("musteri no", "musteri numarasi"), "limit": ("onayli limit", "limit", "onay tutari"),
                  "pb": ("para birimi", "pb"), "urunler": ("urunler", "urun", "izin verilen urunler"), "vade": ("azami vade", "azami vade ay"),
                  "faiz": ("asgari faiz", "asgari faiz orani", "asgari faiz yuzde"), "bitis": ("gecerlilik bitis", "limit vadesi", "gecerlilik tarihi"),
                  "risk": ("mevcut risk", "kullanilan limit", "mevcut kullanim"), "soz_tutar": ("sozlesme tutari", "gks tutari"), "soz_durum": ("sozlesme durumu", "sozlesme", "gks durumu")}
SART_SUTUNLARI = {"onay": ("onay no", "tahsis no", "karar no", "limit no"), "sart": ("sart", "onay sarti", "aciklama"), "tur": ("tur", "sart turu"),
                  "zaman": ("zamanlama", "ne zaman"), "gereken": ("gereken tutar", "istenen tutar", "tutar"), "tesis": ("tesis edilen tutar", "tesis edilen", "alinan tutar"),
                  "durum": ("durum",), "tarih": ("tamamlanma tarihi", "tesis tarihi")}
HEPSI = {"", "*", "tumu", "hepsi", "tum"}
SONUC_RENK = {"Kullandırılabilir": YESIL, "Kontrol edin": SARI, "Engel var": KIRMIZI}


def tl(x) -> str:
    return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def pb_duzelt(x) -> str:
    return metin(x).upper().replace("TRY", "TL") or "TL"


def tamam_mi(x) -> bool:
    return katla(x).startswith(("tamam", "tesis edildi", "alindi", "evet", "var", "imzali", "imzalandi"))


def oku(t_yol: Path, o_yol: Path, s_yol: Path | None) -> tuple[list[dict], dict, dict, list[dict]]:
    uy, talepler, onaylar, sartlar = [], [], {}, defaultdict(list)
    for r in kayitlar(o_yol, ONAY_SUTUNLARI, ("onay", "limit")):
        no, limit = metin(r["onay"]), para(r["limit"])
        if not no:
            continue
        if limit is None or limit <= 0 or no in onaylar:
            uy.append(uyari("Yüksek", "Onay satırı okunamadı", no, f"Onay satırı {r['_satir']}: " + ("onay numarası tekrar ediyor" if no in onaylar else f"limit '{metin(r['limit'])}'") + ". Satır atlandı."))
            continue
        bitis = tarih(r["bitis"])
        if metin(r["bitis"]) and not bitis:
            uy.append(uyari("Yüksek", "Geçerlilik tarihi okunamadı", no, f"'{metin(r['bitis'])}'; geçerlilik kontrolü yapılamadı."))
        onaylar[no] = {"onay": no, "musteri_no": metin(r["musteri_no"]), "limit": limit, "pb": pb_duzelt(r["pb"]),
                       "urunler": {katla(u) for u in re.split(r"[,|+]", metin(r["urunler"])) if katla(u)} - HEPSI, "urunler_ad": metin(r["urunler"]),
                       "vade": para(r["vade"]), "faiz": para(r["faiz"]), "bitis": bitis, "risk": para(r["risk"]) or Decimal(0), "soz_tutar": para(r["soz_tutar"]),
                       "soz_durum": metin(r["soz_durum"]), "soz_var": bool(metin(r["soz_durum"])) or para(r["soz_tutar"]) is not None}
    if s_yol:
        for r in kayitlar(s_yol, SART_SUTUNLARI, ("onay", "sart")):
            if not metin(r["onay"]) or not metin(r["sart"]):
                continue
            if metin(r["onay"]) not in onaylar:
                uy.append(uyari("Bilgi", "Şartın onayı listede yok", r["onay"], f"Şart satırı {r['_satir']}: '{metin(r['sart'])}'."))
                continue
            sartlar[metin(r["onay"])].append({"sart": metin(r["sart"]), "tur": metin(r["tur"]) or "Şart", "sonra": katla(r["zaman"]).startswith(("kullandirim sonrasi", "sonra")),
                                              "zaman": metin(r["zaman"]) or "Kullandırım öncesi", "gereken": para(r["gereken"]), "tesis": para(r["tesis"]),
                                              "tamam": tamam_mi(r["durum"]), "durum": metin(r["durum"]) or "Bekliyor", "tarih": tarih(r["tarih"])})
    for r in kayitlar(t_yol, TALEP_SUTUNLARI, ("no", "onay", "tutar")):
        no = metin(r["no"])
        if not no:
            continue
        tutar = para(r["tutar"])
        if tutar is None or tutar <= 0:
            uy.append(uyari("Yüksek", "Talep tutarı okunamadı", no, f"Talep satırı {r['_satir']}: '{metin(r['tutar'])}'. Talep atlandı."))
            continue
        t = tarih(r["tarih"])
        if metin(r["tarih"]) and not t:
            uy.append(uyari("Orta", "Kullandırım tarihi okunamadı", no, f"'{metin(r['tarih'])}'; kontrol tarihi kullanıldı."))
        talepler.append({"no": no, "musteri_no": metin(r["musteri_no"]), "musteri": metin(r["musteri"]), "sube": metin(r["sube"]), "onay": metin(r["onay"]), "urun": metin(r["urun"]),
                         "tutar": tutar, "pb": pb_duzelt(r["pb"]), "vade": para(r["vade"]), "faiz": para(r["faiz"]), "tarih": t})
    return talepler, onaylar, sartlar, uy


def kontrol(talepler: list[dict], onaylar: dict, sartlar: dict, bugun: date, yaklasan: int = 15) -> dict:
    kullanilan = defaultdict(Decimal)                                      # bu çalıştırmada onay bazında ayrılan tutar
    sonuc, takip = [], []
    for t in talepler:
        b: list[tuple] = []
        o = onaylar.get(t["onay"])
        kt = t["tarih"] or bugun
        if o is None:
            b.append(("Engel", "Onay", "Onay kaydı bulunamadı", f"Onay no '{t['onay']}' onay dosyasında yok.", "Geçerli tahsis kararını kontrol edin."))
            sonuc.append(t | {"bulgular": b, "sonuc": "Engel var", "kullanilabilir": None, "limit": None})
            continue
        bos = o["limit"] - o["risk"] - kullanilan[o["onay"]]
        if o["musteri_no"] and t["musteri_no"] and o["musteri_no"] != t["musteri_no"]:
            b.append(("Engel", "Onay", "Müşteri onayla uyuşmuyor", f"Talep müşterisi {t['musteri_no']}, onay müşterisi {o['musteri_no']}.", "Doğru onay numarasını girin."))
        if o["bitis"]:
            if kt > o["bitis"]:
                b.append(("Engel", "Onay", "Onayın geçerliliği dolmuş", f"Geçerlilik {o['bitis']:%d.%m.%Y}, kullandırım {kt:%d.%m.%Y}.", "Limit yenileme / temdit kararı gerekir."))
            elif (o["bitis"] - kt).days <= yaklasan:
                b.append(("Kontrol", "Onay", "Onayın geçerliliği dolmak üzere", f"Geçerlilik {o['bitis']:%d.%m.%Y} ({(o['bitis'] - kt).days} gün kaldı).", "Yenileme sürecini başlatın."))
        if t["pb"] != o["pb"]:
            b.append(("Engel", "Limit", "Para birimi onaydan farklı", f"Talep {t['pb']}, onay {o['pb']}.", "Onay farklı para biriminde kullanıma izin veriyorsa onay kaydını güncelleyin."))
        elif t["tutar"] > bos:
            b.append(("Engel", "Limit", "Limit yetersiz", f"Onaylı limit {tl(o['limit'])}, mevcut risk {tl(o['risk'])}" + (f", bu listede ayrılan {tl(kullanilan[o['onay']])}" if kullanilan[o["onay"]] else "") +
                      f"; kullanılabilir {tl(max(bos, Decimal(0)))}, talep {tl(t['tutar'])} (aşım {tl(t['tutar'] - bos)}).", "Talep tutarını düşürün veya limit artışı onayı alın."))
        if o["urunler"] and katla(t["urun"]) not in o["urunler"]:
            b.append(("Engel", "Ürün", "Ürün onay kapsamında değil", f"Talep '{t['urun']}', onaydaki ürünler: {o['urunler_ad']}.", "Onay kapsamındaki bir ürün seçin veya ek onay alın."))
        if o["vade"] is not None and t["vade"] is not None and t["vade"] > o["vade"]:
            b.append(("Engel", "Vade", "Vade azami vadeyi aşıyor", f"Talep {t['vade']:f} ay, azami {o['vade']:f} ay.", "Vadeyi kısaltın veya ek onay alın."))
        if o["faiz"] is not None and t["faiz"] is not None and t["faiz"] < o["faiz"]:
            b.append(("Engel", "Fiyatlama", "Faiz asgari oranın altında", f"Talep %{t['faiz']:f}, asgari %{o['faiz']:f}.", "Fiyatlama istisnası için yetkili onayı gerekir."))
        if o["soz_var"]:
            if not tamam_mi(o["soz_durum"]):
                b.append(("Engel", "Sözleşme", "Sözleşme imzalı değil", f"Sözleşme durumu: '{o['soz_durum'] or 'boş'}'.", "Genel kredi sözleşmesini imzalatın ve kaydı güncelleyin."))
            elif o["soz_tutar"] is not None and t["pb"] == o["pb"] and o["risk"] + kullanilan[o["onay"]] + t["tutar"] > o["soz_tutar"]:
                b.append(("Engel", "Sözleşme", "Sözleşme tutarı yetersiz", f"Sözleşme {tl(o['soz_tutar'])}; mevcut risk + talep {tl(o['risk'] + kullanilan[o['onay']] + t['tutar'])}.",
                          "Ek sözleşme / tutar artırımı gerekir."))
        else:
            b.append(("Kontrol", "Sözleşme", "Sözleşme bilgisi girilmemiş", "Onay kaydında sözleşme tutarı ve durumu boş.", "Genel kredi sözleşmesini kontrol edin."))
        for s in sartlar.get(o["onay"], []):
            eksik_tutar = s["gereken"] is not None and (s["tesis"] or Decimal(0)) < s["gereken"]
            if s["sonra"]:
                if not s["tamam"]:
                    if not any(x["s"] is s for x in takip):               # aynı şart, aynı onaya bağlı her talepte yinelenmesin
                        takip.append({"t": t, "s": s})
                    b.append(("Bilgi", s["tur"], "Kullandırım sonrası şart", s["sart"], "Takip listesine alındı."))
            elif not s["tamam"]:
                b.append(("Engel", s["tur"], "Kullandırım öncesi şart tamamlanmamış", f"{s['sart']} (durum: {s['durum']}).", "Şart tamamlanmadan kullandırım yapılamaz."))
            elif eksik_tutar:
                b.append(("Engel", s["tur"], "Teminat tutarı eksik", f"{s['sart']}: istenen {tl(s['gereken'])}, tesis edilen {tl(s['tesis'] or Decimal(0))} (eksik {tl(s['gereken'] - (s['tesis'] or Decimal(0)))}).",
                          "Eksik teminatı tamamlayın."))
            elif s["tarih"] and s["tarih"] > kt:
                b.append(("Kontrol", s["tur"], "Şartın tamamlanma tarihi kullandırımdan sonra", f"{s['sart']}: {s['tarih']:%d.%m.%Y}.", "Tarihi doğrulayın."))
        if kt.weekday() >= 5:
            b.append(("Kontrol", "Tarih", "Kullandırım tarihi hafta sonu", f"{kt:%d.%m.%Y}.", "Valör tarihini doğrulayın."))
        elif kt < bugun:
            b.append(("Kontrol", "Tarih", "Kullandırım tarihi geçmiş", f"{kt:%d.%m.%Y}.", "Geriye valörlü işlem için yetki gerekir."))
        durum = "Engel var" if any(x[0] == "Engel" for x in b) else "Kontrol edin" if any(x[0] == "Kontrol" for x in b) else "Kullandırılabilir"
        if durum != "Engel var":
            kullanilan[o["onay"]] += t["tutar"]
        sira = {"Engel": 0, "Kontrol": 1, "Bilgi": 2}
        sonuc.append(t | {"bulgular": sorted(b, key=lambda x: sira[x[0]]), "sonuc": durum, "kullanilabilir": max(bos, Decimal(0)), "limit": o["limit"]})
    return {"talepler": sonuc, "takip": takip, "bugun": bugun}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

ONEM = {"Engel": "Yüksek", "Kontrol": "Orta", "Bilgi": "Bilgi"}


def rapor_yaz(cikti: Path, r: dict, uyarilar: list[dict]) -> None:
    wb = Workbook()
    ts = wb.active
    ts.title = "Talep Sonuçları"
    _baslik(ts, ["Talep No", "Şube", "Müşteri No", "Müşteri", "Onay No", "Ürün", "Para Birimi", "Talep Tutarı", "Onaylı Limit", "Kullanılabilir Limit", "Sonuç", "Engel", "Kontrol",
                 "Özet", "Onaylayan"], (13, 14, 12, 28, 13, 18, 8, 18, 18, 18, 17, 7, 8, 80, 16))
    for t in r["talepler"]:
        engel, ktr = sum(x[0] == "Engel" for x in t["bulgular"]), sum(x[0] == "Kontrol" for x in t["bulgular"])
        _satir(ts, [t["no"], t["sube"], t["musteri_no"], t["musteri"], t["onay"], t["urun"], t["pb"], t["tutar"], t["limit"], t["kullanilabilir"], t["sonuc"], engel, ktr,
                    "; ".join(x[2] for x in t["bulgular"] if x[0] != "Bilgi"), ""], para_s=(8, 9, 10), dolgu={11: SONUC_RENK[t["sonuc"]], 15: SARI}, sar=(14,))
    ts.auto_filter.ref = f"A1:O{ts.max_row}"

    bl = wb.create_sheet("Bulgular")
    _baslik(bl, ["Talep No", "Müşteri", "Düzey", "Konu", "Bulgu", "Ayrıntı", "Yapılacak", "Giderildi"], (13, 28, 9, 14, 38, 90, 60, 10))
    for t in r["talepler"]:
        for duzey, konu, baslik, ayrinti, oneri in t["bulgular"]:
            _satir(bl, [t["no"], t["musteri"], duzey, konu, baslik, ayrinti, oneri, ""], dolgu={3: RENK[ONEM[duzey]], 8: SARI}, sar=(6, 7))
    bl.auto_filter.ref = f"A1:H{bl.max_row}"

    tk = wb.create_sheet("Şart Takibi")
    _baslik(tk, ["Talep No", "Müşteri", "Onay No", "Şart Türü", "Kullandırım Sonrası Şart", "Durum", "Sorumlu", "Hedef Tarih", "Tamamlandı"], (13, 28, 13, 14, 80, 14, 16, 12, 11))
    for x in r["takip"]:
        _satir(tk, [x["t"]["no"], x["t"]["musteri"], x["t"]["onay"], x["s"]["tur"], x["s"]["sart"], x["s"]["durum"], "", "", ""], dolgu={7: SARI, 8: SARI, 9: SARI}, sar=(5,))
    if not r["takip"]:
        tk.append(["Kullandırım sonrası takip edilecek şart yok."])
    uyari_sayfasi(wb, uyarilar)
    kaydet(wb, cikti)


def calistir(t_yol: Path, o_yol: Path, cikti: Path, bugun: date, s_yol: Path | None = None, yaklasan: int = 15) -> dict:
    talepler, onaylar, sartlar, uy = oku(t_yol, o_yol, s_yol)
    if not talepler:
        raise ValueError(f"{t_yol.name}: okunabilir talep yok")
    r = kontrol(talepler, onaylar, sartlar, bugun, yaklasan)
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, r["uyarilar"])
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Kullandırım taleplerini onay şartları, teminat tesis durumu ve sözleşme bilgileriyle karşılaştırır.")
    p.add_argument("--talepler", type=Path, default=ORNEK / "talepler.csv", help="Talep No, Müşteri No, Müşteri, Şube, Onay No, Ürün, Tutar, Para Birimi, Vade (ay), Faiz Oranı, Kullandırım Tarihi")
    p.add_argument("--onaylar", type=Path, default=ORNEK / "onaylar.csv",
                   help="Onay No, Müşteri No, Onaylı Limit, Para Birimi, Ürünler, Azami Vade, Asgari Faiz, Geçerlilik Bitiş, Mevcut Risk, Sözleşme Tutarı, Sözleşme Durumu")
    p.add_argument("--sartlar", type=Path, help="Onay No, Şart, Tür, Zamanlama (Kullandırım öncesi / sonrası), Gereken Tutar, Tesis Edilen Tutar, Durum, Tamamlanma Tarihi")
    p.add_argument("--yaklasan", type=int, default=15, help="Onay geçerliliği bu kadar gün içinde doluyorsa uyar (varsayılan 15)")
    p.add_argument("--tarih", help="Kontrol tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 09.10.2026)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kullandirim_on_kontrol.xlsx")
    a = p.parse_args(argv)
    ornek = a.talepler == ORNEK / "talepler.csv"
    sart = a.sartlar or (ORNEK / "sartlar.csv" if ornek else None)
    for y in (a.talepler, a.onaylar, sart):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    bugun = tarih(a.tarih) if a.tarih else (date(2026, 10, 9) if ornek else date.today())
    if not bugun:
        print("[X] --tarih GG.AA.YYYY biçiminde olmalı")
        return 2
    try:
        r = calistir(a.talepler, a.onaylar, a.cikti, bugun, sart, a.yaklasan)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    say = {s: sum(t["sonuc"] == s for t in r["talepler"]) for s in SONUC_RENK}
    print(f"[OK] {len(r['talepler'])} talep · kullandırılabilir {say['Kullandırılabilir']} · kontrol edin {say['Kontrol edin']} · engel var {say['Engel var']}")
    for t in r["talepler"]:
        if t["sonuc"] == "Engel var":
            print(f"[!] {t['no']} {t['musteri']}: " + "; ".join(x[2] for x in t["bulgular"] if x[0] == "Engel"))
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
