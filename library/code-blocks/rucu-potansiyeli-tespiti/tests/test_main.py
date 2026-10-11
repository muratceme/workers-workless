import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
RAPOR = date(2026, 10, 1)


def hasar(no="H1", odenen=1000, kusur=100, karsi="Karşı", sigortaci="", tahsil=0, t=date(2026, 6, 1), odeme=date(2026, 7, 1)):
    return {"no": no, "police": "", "sigortali": "", "brans": "", "odenen": D(odenen), "kusur": None if kusur is None else D(kusur), "tahsil": D(tahsil), "karsi": karsi, "sigortaci": sigortaci,
            "karsi_police": "", "hasar": t, "odeme": odeme}


def tek(rapor=RAPOR, secenek=None, **kw):
    return main.degerlendir([hasar(**kw)], rapor, **(secenek or {}))["hasarlar"][0]


class DegerlendirmeTesti(unittest.TestCase):
    def test_tutar_ve_muhatap(self):
        h = tek(odenen=80000, kusur=75)
        self.assertEqual((h["potansiyel"], h["kalan"], h["durum"], h["muhatap"], h["muhatap_tur"]), (D("60000.00"), D("60000.00"), "Rücu edilebilir", "Karşı", "Karşı taraf"))
        h = tek(odenen=40000, sigortaci="X Sigorta", tahsil=15000)
        self.assertEqual((h["kalan"], h["muhatap"], h["muhatap_tur"]), (D("25000.00"), "X Sigorta", "Sigorta şirketi"))
        self.assertEqual(tek(odenen=333, kusur="33.3")["potansiyel"], D("110.89"))

    def test_durumlar(self):
        d = lambda **kw: tek(**kw)["durum"]   # noqa: E731
        self.assertEqual(d(odenen=0), "Ödeme yapılmamış")
        self.assertEqual(d(kusur=None), "Kusur bilgisi eksik")
        self.assertEqual(d(kusur=None, karsi=""), "Rücu yok")
        self.assertEqual(d(kusur=0), "Rücu yok")
        self.assertEqual(d(karsi=""), "Karşı taraf belirsiz")
        self.assertEqual(d(tahsil=1000), "Tamamlandı")
        self.assertEqual(d(tahsil=1500), "Tamamlandı")
        self.assertEqual(tek(secenek={"sirket": "bizim sigorta"}, sigortaci="Bizim Sigorta")["durum"], "Aynı şirket")
        self.assertEqual([tek(secenek={"asgari": D(a)}, kusur=50)["durum"] for a in (500, 501)], ["Rücu edilebilir", "Asgari tutarın altında"])
        r = main.degerlendir([hasar(tahsil=1500)], RAPOR)
        self.assertEqual([u["tur"] for u in r["uyarilar"]], ["Tahsilat potansiyelden büyük"])

    def test_takip_suresi(self):
        h = tek(t=date(2024, 11, 20))
        self.assertEqual((h["takip"], h["kalan_gun"], h["durum"]), (date(2026, 11, 20), 50, "Rücu edilebilir · acil"))
        h = tek(t=date(2024, 8, 2))
        self.assertEqual((h["kalan_gun"], h["durum"]), (-60, "Takip süresi geçmiş"))
        self.assertEqual([tek(t=t)["durum"] for t in (date(2024, 12, 30), date(2024, 12, 31), date(2024, 10, 1), date(2024, 9, 30))],
                         ["Rücu edilebilir · acil", "Rücu edilebilir", "Rücu edilebilir · acil", "Takip süresi geçmiş"])          # 90. gün acil, 0. gün hâlâ süre içinde
        self.assertEqual(tek(t=date(2024, 8, 2), secenek={"baz": "odeme"}, odeme=date(2025, 1, 15))["durum"], "Rücu edilebilir")
        self.assertEqual(tek(t=date(2024, 8, 2), secenek={"takip_yil": 5})["takip"], date(2029, 8, 2))
        self.assertEqual(main.yil_ekle(date(2024, 2, 29), 2), date(2026, 2, 28))
        h = tek(t=None, odeme=None)
        self.assertEqual((h["durum"], h["takip"], "hesaplanamadı" in h["not"]), ("Rücu edilebilir", None, True))

    def test_ozetler(self):
        l = [hasar("H1", 1000, sigortaci="A Sig"), hasar("H2", 3000, sigortaci="a sig", t=date(2024, 11, 20)), hasar("H3", 500, 50), hasar("H4", 900, 0), hasar("H5", 0)]
        r = main.degerlendir(l, RAPOR)
        self.assertEqual([h["no"] for h in r["hasarlar"]], ["H2", "H1", "H3", "H5", "H4"])
        self.assertEqual((r["toplam"], len(r["acik"])), (D("4250.00"), 3))
        self.assertEqual([(m["ad"], m["dosya"], m["tutar"], m["acil"]) for m in r["muhatap"]], [("a sig", 2, D("4000.00"), 1), ("Karşı", 1, D("250.00"), 0)])
        self.assertEqual({d: x["dosya"] for d, x in r["durum"].items()}, {"Rücu edilebilir · acil": 1, "Rücu edilebilir": 2, "Ödeme yapılmamış": 1, "Rücu yok": 1})


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "r.xlsx"
            r = main.calistir(ORNEK / "odenen_hasarlar.csv", cikti, RAPOR)
            d = {h["no"]: (h["durum"], h["kalan"]) for h in r["hasarlar"]}
            self.assertEqual(d["H-301"], ("Rücu edilebilir", D("120000.00")))
            self.assertEqual(d["H-302"], ("Rücu edilebilir", D("60000.00")))
            self.assertEqual(d["H-303"], ("Rücu edilebilir · acil", D("900000.00")))
            self.assertEqual(d["H-304"], ("Takip süresi geçmiş", D("22500.00")))
            self.assertEqual((d["H-305"][0], d["H-306"][0], d["H-307"][0], d["H-308"][0], d["H-310"][0]),
                             ("Kusur bilgisi eksik", "Karşı taraf belirsiz", "Rücu yok", "Tamamlandı", "Ödeme yapılmamış"))
            self.assertEqual((d["H-309"], d["H-311"]), (("Rücu edilebilir", D("25000.00")), ("Rücu edilebilir", D("2000.00"))))
            self.assertEqual((r["toplam"], len(r["acik"]), r["hasarlar"][0]["no"]), (D("1107000.00"), 5, "H-303"))
            self.assertEqual([(m["ad"], m["tutar"]) for m in r["muhatap"]][:2], [("Hayali Sigorta", D("925000.00")), ("Kurgu Sigorta", D("120000.00"))])
            self.assertEqual(r["uyarilar"], [])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Rücu Dosyaları", "Muhatap Bazında", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "h.csv"
            y.write_text("Hasar No;Hasar Tarihi;Ödenen Tazminat;Karşı Taraf;Karşı Taraf Kusur %\nH1;01.06.2026;1.000,00;X;%75\nH1;;5;;\nH2;;çok;;\nH3;dün;100;Y;150\nH4;;-5;;\n",
                         encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual([(h["no"], h["odenen"], h["kusur"], h["hasar"]) for h in l], [("H1", D("1000.00"), D(75), date(2026, 6, 1)), ("H3", D(100), None, None)])
            self.assertEqual(sorted(u["tur"] for u in uy), ["Kusur oranı okunamadı", "Mükerrer hasar dosyası", "Tarih okunamadı", "Ödenen tutar okunamadı", "Ödenen tutar okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--hasarlar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "01.10.2026", "--baz", "odeme", "--asgari", "500", "--sirket", "Bizim"]), 0)
            self.assertEqual(main.main(arg + ["--tarih", "dün"]), 2)
            self.assertEqual(main.main(arg + ["--takip-yil", "0"]), 2)
            self.assertEqual(main.main(["--hasarlar", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Hasar No;Ödenen Tazminat\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
