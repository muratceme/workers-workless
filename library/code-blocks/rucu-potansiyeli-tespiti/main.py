"""
Rücu Potansiyeli Tespiti — Workers / Workless kod bloğu
Hasar Yönetimi › Rücu Uzmanı

Ödenen hasar dosyalarını kusur oranı, karşı taraf ve sigorta bilgisine göre tarar; rücu edilebilir dosyaları ve tutarlarını listeler.
  Rücu potansiyeli = ödenen tazminat × karşı taraf kusur oranı %          (sigortacı, ödediği tutarla sınırlı olarak sigortalının yerine geçer)
  Kalan tutar      = rücu potansiyeli − tahsil edilen rücu
  Muhatap          = karşı tarafın sorumluluk sigortacısı biliniyorsa o; bilinmiyorsa karşı tarafın kendisi
  Takip tarihi     = hasar tarihi (veya --baz odeme ile ödeme tarihi) + --takip-yil (varsayılan 2) yıl;  kalan gün ≤ --acil-gun (90) ise "acil"
Takip tarihi bir iç takip süresidir, hukuki zamanaşımı hesabı DEĞİLDİR: zamanaşımı olayın hukuki niteliğine göre değişir; hukuk birimine danışın.
Kod kusur tespiti yapmaz; kusur oranını dosyadan okur.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"no": ("hasar no", "dosya no", "hasar dosya no"), "police": ("police no",), "sigortali": ("sigortali", "sigortali adi"), "brans": ("brans", "hasar turu"),
         "hasar": ("hasar tarihi",), "odeme": ("odeme tarihi", "tazminat odeme tarihi"), "odenen": ("odenen tazminat", "odenen", "odenen tutar"),
         "karsi": ("karsi taraf", "sorumlu", "ucuncu kisi"), "kusur": ("karsi taraf kusur", "karsi taraf kusur orani", "karsi kusur", "kusur orani"),
         "sigortaci": ("karsi taraf sigortacisi", "karsi sigortaci", "karsi sigorta sirketi"), "karsi_police": ("karsi police no", "karsi taraf police no"),
         "tahsil": ("tahsil edilen rucu", "tahsil edilen", "rucu tahsilati")}
Z = Decimal(0)
SIRA = ["Rücu edilebilir · acil", "Rücu edilebilir", "Takip süresi geçmiş", "Karşı taraf belirsiz", "Kusur bilgisi eksik", "Asgari tutarın altında", "Aynı şirket", "Ödeme yapılmamış",
        "Rücu yok", "Tamamlandı"]
ACIKLAMA = {"Rücu edilebilir · acil": "Takip süresinin dolmasına az kaldı; dosyayı öne alın.", "Rücu edilebilir": "Muhataba rücu talebi gönderilebilir.",
            "Takip süresi geçmiş": "İç takip süresi dolmuş. Hukuki zamanaşımı ayrıca değerlendirilmeli; hukuk birimine danışın.",
            "Karşı taraf belirsiz": "Karşı taraf kusurlu görünüyor ama kimliği girilmemiş; tutanak ve dosyadan araştırın.",
            "Kusur bilgisi eksik": "Karşı taraf var, kusur oranı girilmemiş; eksper raporu veya tutanaktan tamamlayın.",
            "Asgari tutarın altında": "Kalan tutar asgari rücu tutarının altında.", "Aynı şirket": "Karşı tarafın sigortacısı kendi şirketiniz; iç mahsup usulünüze göre işlem yapın.",
            "Ödeme yapılmamış": "Tazminat ödenmeden halefiyet doğmaz; ödeme sonrası yeniden değerlendirin.", "Rücu yok": "Karşı tarafın kusuru yok veya karşı taraf yok.",
            "Tamamlandı": "Rücu potansiyelinin tamamı tahsil edilmiş."}


def yil_ekle(t: date, yil: int) -> date:
    try:
        return t.replace(year=t.year + yil)
    except ValueError:
        return t.replace(year=t.year + yil, day=28)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("no", "odenen")):
        no = metin(r["no"])
        if not no:
            continue
        kim = f"{no} (satır {r['_satir']})"
        if no in gorulen:
            uy.append(uyari("Yüksek", "Mükerrer hasar dosyası", kim, "Hasar numarası tekrar ediyor; ilk satır kullanıldı."))
            continue
        odenen, kusur, tahsil = para(r["odenen"]), para(r["kusur"]), para(r["tahsil"])
        if odenen is None and metin(r["odenen"]) or (odenen is not None and odenen < 0):
            uy.append(uyari("Yüksek", "Ödenen tutar okunamadı", kim, f"'{metin(r['odenen'])}'. Dosya değerlendirilmedi."))
            continue
        if metin(r["kusur"]) and (kusur is None or not 0 <= kusur <= 100):
            uy.append(uyari("Yüksek", "Kusur oranı okunamadı", kim, f"'{metin(r['kusur'])}'. 0-100 arası yüzde olmalı; kusur bilgisi eksik sayıldı."))
            kusur = None
        h = {"no": no, "police": metin(r["police"]), "sigortali": metin(r["sigortali"]), "brans": metin(r["brans"]), "odenen": odenen or Z, "kusur": kusur, "tahsil": tahsil or Z,
             "karsi": metin(r["karsi"]), "sigortaci": metin(r["sigortaci"]), "karsi_police": metin(r["karsi_police"])}
        for a, ad in (("hasar", "Hasar tarihi"), ("odeme", "Ödeme tarihi")):
            h[a] = tarih(r[a])
            if metin(r[a]) and not h[a]:
                uy.append(uyari("Orta", "Tarih okunamadı", kim, f"{ad} '{metin(r[a])}'. Takip tarihi hesaplanamayabilir."))
        gorulen.add(no)
        l.append(h)
    return l, uy


def degerlendir(hasarlar: list[dict], rapor: date, takip_yil: int = 2, acil_gun: int = 90, asgari: Decimal = Z, baz: str = "hasar", sirket: str = "") -> dict:
    uy = []
    for h in hasarlar:
        b = h[baz] or (h["hasar"] if baz == "odeme" else None)
        h["takip"] = yil_ekle(b, takip_yil) if b else None
        h["kalan_gun"] = (h["takip"] - rapor).days if h["takip"] else None
        h["potansiyel"] = yuvarla(h["odenen"] * h["kusur"] / 100) if h["kusur"] is not None else None
        h["kalan"] = max(h["potansiyel"] - h["tahsil"], Z) if h["potansiyel"] is not None else None
        h["muhatap"] = h["sigortaci"] or h["karsi"]
        h["muhatap_tur"] = "Sigorta şirketi" if h["sigortaci"] else "Karşı taraf" if h["karsi"] else ""
        if h["odenen"] <= 0:
            d = "Ödeme yapılmamış"
        elif h["kusur"] is None:
            d = "Kusur bilgisi eksik" if h["karsi"] or h["sigortaci"] else "Rücu yok"
        elif h["kusur"] == 0:
            d = "Rücu yok"
        elif not h["muhatap"]:
            d = "Karşı taraf belirsiz"
        elif h["kalan"] == 0:
            d = "Tamamlandı"
        elif sirket and katla(h["sigortaci"]) == katla(sirket):
            d = "Aynı şirket"
        elif h["kalan"] < asgari:
            d = "Asgari tutarın altında"
        elif h["kalan_gun"] is not None and h["kalan_gun"] < 0:
            d = "Takip süresi geçmiş"
        elif h["kalan_gun"] is not None and h["kalan_gun"] <= acil_gun:
            d = "Rücu edilebilir · acil"
        else:
            d = "Rücu edilebilir"
        h["durum"], h["not"] = d, ACIKLAMA[d]
        if d.startswith("Rücu edilebilir") and h["kalan_gun"] is None:
            h["not"] += " Hasar tarihi olmadığı için takip tarihi hesaplanamadı."
        if h["potansiyel"] is not None and h["tahsil"] > h["potansiyel"]:
            uy.append(uyari("Orta", "Tahsilat potansiyelden büyük", h["no"], f"Tahsil edilen {h['tahsil']:f} > rücu potansiyeli {h['potansiyel']:f}. Kusur oranını veya tahsilatı kontrol edin."))
    hasarlar.sort(key=lambda h: (SIRA.index(h["durum"]), -(h["kalan"] or Z), h["no"]))
    acik = [h for h in hasarlar if h["durum"].startswith("Rücu edilebilir")]
    muhatap = {}
    for h in acik:
        m = muhatap.setdefault(katla(h["muhatap"]), {"ad": h["muhatap"], "tur": h["muhatap_tur"], "dosya": 0, "tutar": Z, "acil": 0})
        m["dosya"] += 1
        m["tutar"] += h["kalan"]
        m["acil"] += h["durum"].endswith("acil")
    durum = {}
    for h in hasarlar:
        x = durum.setdefault(h["durum"], {"dosya": 0, "odenen": Z, "kalan": Z})
        x["dosya"] += 1
        x["odenen"] += h["odenen"]
        x["kalan"] += h["kalan"] or Z
    return {"hasarlar": hasarlar, "acik": acik, "toplam": sum((h["kalan"] for h in acik), Z), "muhatap": sorted(muhatap.values(), key=lambda m: -m["tutar"]),
            "durum": {d: durum[d] for d in SIRA if d in durum}, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, rapor: date) -> None:
    wb = Workbook()
    renk = {"Rücu edilebilir · acil": KIRMIZI, "Rücu edilebilir": YESIL, "Takip süresi geçmiş": KIRMIZI, "Karşı taraf belirsiz": SARI, "Kusur bilgisi eksik": SARI}
    ws = wb.active
    ws.title = "Özet"
    ws.append([f"Rapor tarihi {rapor:%d.%m.%Y}"])
    ws["A1"].font = Font(bold=True)
    _baslik(ws, ["Durum", "Dosya", "Ödenen Tazminat", "Kalan Rücu Tutarı"], (30, 9, 18, 18), "A3")
    for d, x in r["durum"].items():
        _satir(ws, [d, x["dosya"], x["odenen"], x["kalan"]], para_s=(3, 4), dolgu={1: renk.get(d)})
    ws.append([])
    _satir(ws, ["Rücu edilebilir toplam", len(r["acik"]), "", r["toplam"]], para_s=(4,))
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.append([])
    ws.append(["Takip tarihi iç takip süresidir; hukuki zamanaşımı hesabı değildir. Kusur oranı dosyadan okunur, kod kusur tespiti yapmaz."])
    ws = wb.create_sheet("Rücu Dosyaları")
    _baslik(ws, ["Hasar No", "Poliçe No", "Sigortalı", "Branş", "Hasar Tarihi", "Ödenen Tazminat", "Karşı Kusur %", "Rücu Potansiyeli", "Tahsil Edilen", "Kalan Tutar", "Muhatap",
                 "Muhatap Türü", "Karşı Poliçe No", "Takip Tarihi", "Kalan Gün", "Durum", "Not", "Yapılan İşlem (siz doldurun)"],
            (13, 13, 26, 14, 12, 16, 9, 16, 14, 16, 26, 14, 14, 12, 9, 22, 60, 26))
    for h in r["hasarlar"]:
        _satir(ws, [h["no"], h["police"], h["sigortali"], h["brans"], h["hasar"], h["odenen"], h["kusur"], h["potansiyel"], h["tahsil"] or None, h["kalan"], h["muhatap"], h["muhatap_tur"],
                    h["karsi_police"], h["takip"], h["kalan_gun"], h["durum"], h["not"], ""], para_s=(6, 8, 9, 10), tarih_s=(5, 14), dolgu={16: renk.get(h["durum"])}, sar=(17,))
    ws = wb.create_sheet("Muhatap Bazında")
    _baslik(ws, ["Muhatap", "Tür", "Dosya", "Acil Dosya", "Kalan Rücu Tutarı"], (32, 16, 9, 10, 18))
    for m in r["muhatap"]:
        _satir(ws, [m["ad"], m["tur"], m["dosya"], m["acil"] or None, m["tutar"]], para_s=(5,), dolgu={4: KIRMIZI if m["acil"] else None})
    uyari_sayfasi(wb, sirala(r["uyarilar"]))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, rapor: date, **secenek) -> dict:
    hasarlar, uy = oku(yol)
    if not hasarlar:
        raise ValueError(f"{yol.name}: değerlendirilecek hasar dosyası yok.")
    r = degerlendir(hasarlar, rapor, **secenek)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, rapor)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Ödenen hasarlarda rücu edilebilir dosyaları ve tutarlarını belirler.")
    p.add_argument("--hasarlar", type=Path, default=ORNEK / "odenen_hasarlar.csv",
                   help="Hasar No, Poliçe No, Sigortalı, Branş, Hasar Tarihi, Ödeme Tarihi, Ödenen Tazminat, Karşı Taraf, Karşı Taraf Kusur %%, Karşı Taraf Sigortacısı, Tahsil Edilen Rücu")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 01.10.2026)")
    p.add_argument("--takip-yil", type=int, default=2, help="İç takip süresi, yıl (varsayılan 2; hukuki zamanaşımı değildir)")
    p.add_argument("--baz", choices=("hasar", "odeme"), default="hasar", help="Takip süresinin başlangıcı: hasar tarihi veya ödeme tarihi (varsayılan hasar)")
    p.add_argument("--acil-gun", type=int, default=90, help="Takip tarihine bu kadar gün veya daha az kaldıysa acil (varsayılan 90)")
    p.add_argument("--asgari", default="0", help="Asgari rücu tutarı; altındaki dosyalar ayrı gösterilir (varsayılan 0)")
    p.add_argument("--sirket", default="", help="Kendi sigorta şirketinizin adı (karşı sigortacı aynıysa ayrı gösterilir)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "rucu_potansiyeli.xlsx")
    a = p.parse_args(argv)
    if not a.hasarlar.exists():
        print(f"[X] Dosya bulunamadı: {a.hasarlar}")
        return 1
    rapor = tarih(a.tarih) if a.tarih else (date(2026, 10, 1) if a.hasarlar == ORNEK / "odenen_hasarlar.csv" else date.today())
    asgari = para(a.asgari)
    if not rapor or asgari is None or asgari < 0 or a.takip_yil < 1 or a.acil_gun < 0:
        print("[X] --tarih GG.AA.YYYY biçiminde; --asgari sıfır veya pozitif; --takip-yil en az 1 olmalı")
        return 2
    try:
        r = calistir(a.hasarlar, a.cikti, rapor, takip_yil=a.takip_yil, acil_gun=a.acil_gun, asgari=asgari, baz=a.baz, sirket=a.sirket)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(r['hasarlar'])} dosya · {len(r['acik'])} rücu edilebilir · kalan tutar {r['toplam']:,.2f}")
    for h in r["hasarlar"]:
        if h["durum"] in ("Rücu edilebilir · acil", "Takip süresi geçmiş"):
            print(f"[!] {h['no']} · {h['durum']}: {h['kalan']:,.2f} · takip tarihi {h['takip']:%d.%m.%Y} ({h['kalan_gun']} gün)")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
