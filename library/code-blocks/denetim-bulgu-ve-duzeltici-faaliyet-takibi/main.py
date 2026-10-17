"""
Denetim Bulgu ve Düzeltici Faaliyet Takibi — Gerisi Bende kod bloğu
Sürdürülebilirlik ve Sosyal Uygunluk › Sosyal Uygunluk Uzmanı

Sosyal uygunluk ve müşteri denetimlerinin bulgularını düzeltici faaliyet planı (CAP), sorumlu ve termin bazında takip eder.
  Açık bulgu     = durumu "tamamlandı", "kapandı" veya "doğrulandı" ile başlamayan bulgu
  Takip durumu (açık)   : Planı eksik (düzeltici faaliyet, sorumlu veya termin boş) · Gecikmiş (termin < rapor tarihi) ·
                          Yaklaşıyor (termine --yaklasan, varsayılan 14 gün veya daha az kaldı) · Zamanında
  Takip durumu (kapalı) : Kanıt eksik (kanıt "Yok" veya boş) · Geç tamamlandı (tamamlanma > termin) · Tamamlandı
  Tamamlanma %   = kapalı bulgu ÷ toplam bulgu × 100
Kod denetim standardını yorumlamaz, bulgu derecesini değiştirmez ve denetim notu hesaplamaz; yalnız sizin listenizi takip eder.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"no": ("bulgu no", "no", "kayit no"), "denetim": ("denetim", "denetim turu", "standart"), "dtarih": ("denetim tarihi",), "tesis": ("tesis", "fabrika", "lokasyon"),
         "konu": ("konu", "alan", "baslik", "performans alani"), "bulgu": ("bulgu", "uygunsuzluk", "bulgu aciklamasi"), "derece": ("derece", "onem", "sinif"),
         "df": ("duzeltici faaliyet", "aksiyon", "faaliyet", "cap"), "sorumlu": ("sorumlu",), "termin": ("termin", "termin tarihi", "hedef tarih"), "durum": ("durum",),
         "tamam": ("tamamlanma tarihi", "kapanis tarihi"), "kanit": ("kanit", "kanit belgesi", "dokuman")}
KAPALI = ("tamam", "kapa", "dogrula")
DERECE = {"kritik": 0, "sifir tolerans": 0, "major": 1, "buyuk": 1, "minor": 2, "kucuk": 2, "gozlem": 3}
SIRA = ["Gecikmiş", "Planı eksik", "Yaklaşıyor", "Zamanında", "Kanıt eksik", "Geç tamamlandı", "Tamamlandı"]
YOK = ("", "yok", "hayir", "h", "-")


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("bulgu", "durum")):
        if not metin(r["bulgu"]) and not metin(r["no"]):
            continue
        no = metin(r["no"]) or f"satır {r['_satir']}"
        if no in gorulen:
            uy.append(uyari("Yüksek", "Mükerrer bulgu numarası", no, f"{yol.name} satır {r['_satir']}: ilk satır kullanıldı."))
            continue
        gorulen.add(no)
        b = {"no": no, "denetim": metin(r["denetim"]) or "Belirtilmemiş", "tesis": metin(r["tesis"]), "konu": metin(r["konu"]) or "Belirtilmemiş", "bulgu": metin(r["bulgu"]),
             "derece": metin(r["derece"]) or "Belirtilmemiş", "df": metin(r["df"]), "sorumlu": metin(r["sorumlu"]), "durum": metin(r["durum"]) or "Açık",
             "kanit": katla(r["kanit"]) not in YOK}
        b["kapali"] = katla(b["durum"]).startswith(KAPALI)
        b["agirlik"] = next((v for k, v in DERECE.items() if katla(b["derece"]).startswith(k)), 4)
        for a, ad in (("dtarih", "Denetim tarihi"), ("termin", "Termin"), ("tamam", "Tamamlanma tarihi")):
            b[a] = tarih(r[a])
            if metin(r[a]) and not b[a]:
                uy.append(uyari("Orta", "Tarih okunamadı", no, f"{ad} '{metin(r[a])}'. Boş sayıldı."))
        l.append(b)
    return l, uy


def takip(bulgular: list[dict], rapor: date, yaklasan: int = 14) -> dict:
    if not bulgular:
        raise ValueError("Takip edilecek bulgu yok.")
    for b in bulgular:
        notlar = []
        b["kalan"] = (b["termin"] - rapor).days if b["termin"] else None
        if b["kapali"]:
            gec = b["tamam"] and b["termin"] and b["tamam"] > b["termin"]
            if not b["kanit"]:
                notlar.append("Kapanış kanıtı (belge, fotoğraf, kayıt) eklenmemiş.")
            if gec:
                notlar.append(f"Terminden {(b['tamam'] - b['termin']).days} gün sonra tamamlandı.")
            if not b["tamam"]:
                notlar.append("Tamamlanma tarihi girilmemiş.")
            b["takip"] = "Kanıt eksik" if not b["kanit"] else "Geç tamamlandı" if gec else "Tamamlandı"
            b["gecikme"] = None
        else:
            eksik = [ad for ad, v in (("düzeltici faaliyet", b["df"]), ("sorumlu", b["sorumlu"]), ("termin", b["termin"])) if not v]
            if eksik:
                notlar.append("Eksik: " + ", ".join(eksik) + ".")
            b["gecikme"] = -b["kalan"] if b["kalan"] is not None and b["kalan"] < 0 else None
            if b["gecikme"]:
                notlar.append(f"Termin {b['gecikme']} gün geçti.")
            elif b["kalan"] is not None and b["kalan"] <= yaklasan:
                notlar.append(f"Termine {b['kalan']} gün kaldı.")
            b["takip"] = "Gecikmiş" if b["gecikme"] else "Planı eksik" if eksik else "Yaklaşıyor" if b["kalan"] <= yaklasan else "Zamanında"
        b["not"] = " ".join(notlar)
    bulgular.sort(key=lambda b: (SIRA.index(b["takip"]), b["agirlik"], -(b["gecikme"] or 0), b["kalan"] if b["kalan"] is not None else 9999, b["no"]))

    def grupla(alan: str, bos: str = "Belirtilmemiş") -> list[dict]:
        g = {}
        for b in bulgular:
            x = g.setdefault(b[alan] or bos, {"ad": b[alan] or bos, "toplam": 0, "kapali": 0, "acik": 0, "gecikmis": 0, "yaklasan": 0, "plansiz": 0, "kanitsiz": 0, "kritik_acik": 0})
            x["toplam"] += 1
            x["kapali" if b["kapali"] else "acik"] += 1
            x["gecikmis"] += b["takip"] == "Gecikmiş"
            x["yaklasan"] += b["takip"] == "Yaklaşıyor"
            x["plansiz"] += b["takip"] == "Planı eksik"
            x["kanitsiz"] += b["takip"] == "Kanıt eksik"
            x["kritik_acik"] += (not b["kapali"]) and b["agirlik"] == 0
        for x in g.values():
            x["oran"] = yuvarla(Decimal(x["kapali"]) / x["toplam"] * 100)
        return sorted(g.values(), key=lambda x: (-x["gecikmis"], -x["acik"], x["ad"]))

    say = {s: sum(b["takip"] == s for b in bulgular) for s in SIRA}
    kapali = sum(b["kapali"] for b in bulgular)
    return {"bulgular": bulgular, "denetim": grupla("denetim"), "sorumlu": grupla("sorumlu", "Atanmamış"), "konu": grupla("konu"), "derece": grupla("derece"), "say": say,
            "toplam": len(bulgular), "kapali": kapali, "acik": len(bulgular) - kapali, "oran": yuvarla(Decimal(kapali) / len(bulgular) * 100)}


def rapor_yaz(cikti: Path, r: dict, rapor: date, uy: list[dict]) -> None:
    wb = Workbook()
    renk = {"Gecikmiş": KIRMIZI, "Planı eksik": KIRMIZI, "Yaklaşıyor": SARI, "Zamanında": None, "Kanıt eksik": SARI, "Geç tamamlandı": MAVI, "Tamamlandı": YESIL}
    ws = wb.active
    ws.title = "Takip Listesi"
    ws.append([f"Rapor tarihi {rapor:%d.%m.%Y} · {r['toplam']} bulgu · {r['acik']} açık · {r['say']['Gecikmiş']} gecikmiş · tamamlanma %{r['oran']}"])
    ws["A1"].font = Font(bold=True)
    _baslik(ws, ["Takip Durumu", "Bulgu No", "Denetim", "Tesis", "Konu", "Derece", "Bulgu", "Düzeltici Faaliyet", "Sorumlu", "Termin", "Kalan Gün", "Durum", "Tamamlanma", "Kanıt", "Not",
                 "Güncelleme (siz doldurun)"], (15, 10, 18, 14, 22, 10, 44, 44, 18, 12, 9, 14, 12, 8, 44, 30), "A3")
    for b in r["bulgular"]:
        _satir(ws, [b["takip"], b["no"], b["denetim"], b["tesis"], b["konu"], b["derece"], b["bulgu"], b["df"], b["sorumlu"], b["termin"], None if b["kapali"] else b["kalan"], b["durum"], b["tamam"],
                    "Var" if b["kanit"] else "Yok" if b["kapali"] else "", b["not"], ""], tarih_s=(10, 13), dolgu={1: renk[b["takip"]], 6: KIRMIZI if b["agirlik"] == 0 and not b["kapali"] else None},
               sar=(7, 8, 15))
    for ad, veri, kol in (("Denetim Bazında", r["denetim"], "Denetim"), ("Sorumlu Bazında", r["sorumlu"], "Sorumlu"), ("Konu Bazında", r["konu"], "Konu"), ("Derece Bazında", r["derece"], "Derece")):
        ws = wb.create_sheet(ad)
        _baslik(ws, [kol, "Toplam Bulgu", "Kapalı", "Açık", "Gecikmiş", "Yaklaşıyor", "Planı Eksik", "Kanıt Eksik", "Açık Kritik", "Tamamlanma %"], (30, 9, 8, 8, 9, 10, 10, 10, 10, 12))
        for x in veri:
            _satir(ws, [x["ad"], x["toplam"], x["kapali"], x["acik"], x["gecikmis"] or None, x["yaklasan"] or None, x["plansiz"] or None, x["kanitsiz"] or None, x["kritik_acik"] or None, x["oran"]],
                   para_s=(10,), dolgu={5: KIRMIZI if x["gecikmis"] else None, 9: KIRMIZI if x["kritik_acik"] else None})
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, rapor: date, yaklasan: int = 14) -> dict:
    bulgular, uy = oku(yol)
    r = takip(bulgular, rapor, yaklasan)
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, rapor, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Denetim bulgularını düzeltici faaliyet, sorumlu ve termin bazında takip eder.")
    p.add_argument("--bulgular", type=Path, default=ORNEK / "bulgular.csv",
                   help="Bulgu No, Denetim, Denetim Tarihi, Tesis, Konu, Bulgu, Derece, Düzeltici Faaliyet, Sorumlu, Termin, Durum, Tamamlanma Tarihi, Kanıt")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 12.10.2026)")
    p.add_argument("--yaklasan", type=int, default=14, help="Termine bu kadar gün veya daha az kalan bulgular 'Yaklaşıyor' sayılır (varsayılan 14)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "duzeltici_faaliyet_takibi.xlsx")
    a = p.parse_args(argv)
    if not a.bulgular.exists():
        print(f"[X] Dosya bulunamadı: {a.bulgular}")
        return 1
    rapor = tarih(a.tarih) if a.tarih else (date(2026, 10, 12) if a.bulgular == ORNEK / "bulgular.csv" else date.today())
    if not rapor or a.yaklasan < 0:
        print("[X] --tarih GG.AA.YYYY biçiminde; --yaklasan sıfır veya pozitif olmalı")
        return 2
    try:
        r = calistir(a.bulgular, a.cikti, rapor, a.yaklasan)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {r['toplam']} bulgu · {r['acik']} açık · tamamlanma %{r['oran']} · " + " · ".join(f"{s}: {n}" for s, n in r["say"].items() if n))
    for b in r["bulgular"]:
        if b["takip"] in ("Gecikmiş", "Planı eksik"):
            print(f"[!] {b['no']} · {b['takip']} · {b['derece']} · {b['sorumlu'] or 'sorumlu yok'}: {b['not']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
