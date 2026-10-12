import sys
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from decimal import Decimal as D  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
RAPOR = date(2026, 10, 12)


def b(no="B1", termin=10, kapali=False, df="Yapılacak", sorumlu="S1", kanit=False, tamam=None, derece="Majör", denetim="BSCI", konu="İSG"):
    t = lambda g: None if g is None else RAPOR + timedelta(days=g)   # noqa: E731
    return {"no": no, "denetim": denetim, "tesis": "", "konu": konu, "bulgu": "x", "derece": derece, "df": df, "sorumlu": sorumlu, "durum": "Tamamlandı" if kapali else "Açık",
            "kanit": kanit, "kapali": kapali, "agirlik": {"Kritik": 0, "Majör": 1, "Minör": 2}.get(derece, 4), "dtarih": None, "termin": t(termin), "tamam": t(tamam)}


def tek(**kw):
    return main.takip([b(**kw)], RAPOR)["bulgular"][0]


class TakipTesti(unittest.TestCase):
    def test_acik_bulgular(self):
        self.assertEqual([tek(termin=g)["takip"] for g in (-1, 0, 14, 15)], ["Gecikmiş", "Yaklaşıyor", "Yaklaşıyor", "Zamanında"])
        x = tek(termin=-12)
        self.assertEqual((x["gecikme"], x["not"]), (12, "Termin 12 gün geçti."))
        self.assertEqual(tek(termin=3)["not"], "Termine 3 gün kaldı.")
        x = tek(termin=None, df="", sorumlu="")
        self.assertEqual((x["takip"], x["not"]), ("Planı eksik", "Eksik: düzeltici faaliyet, sorumlu, termin."))
        self.assertEqual(tek(termin=-5, sorumlu="")["takip"], "Gecikmiş")                         # gecikme, plan eksikliğinden önce gelir
        self.assertEqual(main.takip([b(termin=20)], RAPOR, yaklasan=30)["bulgular"][0]["takip"], "Yaklaşıyor")

    def test_kapali_bulgular(self):
        self.assertEqual(tek(kapali=True, kanit=True, termin=-10, tamam=-12)["takip"], "Tamamlandı")
        self.assertEqual(tek(kapali=True, kanit=True, termin=-10, tamam=-10)["takip"], "Tamamlandı")
        x = tek(kapali=True, kanit=True, termin=-10, tamam=-3)
        self.assertEqual((x["takip"], x["not"]), ("Geç tamamlandı", "Terminden 7 gün sonra tamamlandı."))
        x = tek(kapali=True, kanit=False, termin=-10, tamam=-3)
        self.assertEqual((x["takip"], "kanıtı" in x["not"], "7 gün sonra" in x["not"]), ("Kanıt eksik", True, True))
        self.assertIn("Tamamlanma tarihi girilmemiş", tek(kapali=True, kanit=True)["not"])

    def test_siralama_ve_ozetler(self):
        l = [b("B1", 30), b("B2", -5, derece="Minör"), b("B3", -2, derece="Kritik", sorumlu="S2"), b("B4", None, df=""), b("B5", -30, kapali=True, kanit=True, tamam=-31, denetim="Sedex"),
             b("B6", 5, sorumlu="", denetim="Sedex")]
        r = main.takip(l, RAPOR)
        self.assertEqual([x["no"] for x in r["bulgular"]], ["B3", "B2", "B6", "B4", "B1", "B5"])     # gecikmişlerde önce kritik
        self.assertEqual((r["toplam"], r["acik"], r["kapali"], r["oran"], r["say"]["Gecikmiş"], r["say"]["Planı eksik"]), (6, 5, 1, D("16.67"), 2, 2))
        d = {x["ad"]: (x["toplam"], x["kapali"], x["gecikmis"], x["plansiz"], x["kritik_acik"], x["oran"]) for x in r["denetim"]}
        self.assertEqual(d, {"BSCI": (4, 0, 2, 1, 1, D("0.00")), "Sedex": (2, 1, 0, 1, 0, D("50.00"))})
        self.assertEqual({x["ad"]: x["acik"] for x in r["sorumlu"]}, {"S1": 3, "S2": 1, "Atanmamış": 1})
        with self.assertRaises(ValueError):
            main.takip([], RAPOR)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "d.xlsx"
            r = main.calistir(ORNEK / "bulgular.csv", cikti, RAPOR)
            self.assertEqual({x["no"]: x["takip"] for x in r["bulgular"]},
                             {"B-01": "Tamamlandı", "B-02": "Gecikmiş", "B-03": "Yaklaşıyor", "B-04": "Planı eksik", "B-05": "Zamanında", "B-06": "Kanıt eksik", "B-07": "Yaklaşıyor", "B-08": "Gecikmiş"})
            self.assertEqual(([x["no"] for x in r["bulgular"]][:3], r["acik"], r["oran"]), (["B-08", "B-02", "B-04"], 6, D("25.00")))
            self.assertEqual({x["no"]: x["gecikme"] for x in r["bulgular"] if x["gecikme"]}, {"B-08": 27, "B-02": 12})
            self.assertEqual({x["ad"]: (x["toplam"], x["acik"], x["gecikmis"]) for x in r["denetim"]}, {"BSCI": (4, 3, 2), "Sedex SMETA": (2, 2, 0), "Müşteri denetimi": (2, 1, 0)})
            self.assertEqual(next(x for x in r["sorumlu"] if x["ad"] == "İK Müdürü")["gecikmis"], 2)
            self.assertEqual(r["uyarilar"], [])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Takip Listesi", "Denetim Bazında", "Sorumlu Bazında", "Konu Bazında", "Derece Bazında", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "b.csv"
            y.write_text("Bulgu No;Bulgu;Derece;Termin;Durum;Kanıt\nB1;x;KRİTİK;yakında;DOĞRULANDI;Fotoğraf\nB1;y;;;;\nB2;z;Büyük;01.11.2026;;hayır\n;;;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual([(x["no"], x["kapali"], x["agirlik"], x["kanit"], x["termin"], x["durum"]) for x in l],
                             [("B1", True, 0, True, None, "DOĞRULANDI"), ("B2", False, 1, False, date(2026, 11, 1), "Açık")])
            self.assertEqual(sorted(u["tur"] for u in uy), ["Mükerrer bulgu numarası", "Tarih okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--bulgular", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "12.10.2026", "--yaklasan", "30"]), 0)
            self.assertEqual(main.main(arg + ["--tarih", "dün"]), 2)
            self.assertEqual(main.main(["--bulgular", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Bulgu;Durum\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
