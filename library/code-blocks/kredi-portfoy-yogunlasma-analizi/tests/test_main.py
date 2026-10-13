import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def k(musteri, risk, sektor="S", grup=None, teminat="T"):
    return {"musteri": musteri, "grup": grup or musteri, "sektor": sektor, "teminat": teminat, "risk": D(str(risk))}


def lim(boyut, deger, limit):
    return {"boyut": boyut, "deger": deger, "limit": D(str(limit))}


class AnalizTesti(unittest.TestCase):
    def test_hhi(self):
        b = main.analiz([k("a", 100)])["boyutlar"]["Sektör"]
        self.assertEqual((b["hhi"], b["esdeger"], b["en_buyuk"], b["ilk_pay"], b["adet"]), (D("10000.00"), D("1.00"), D("100.00"), D("100.00"), 1))
        b = main.analiz([k(str(i), 25, f"S{i}") for i in range(4)], ilk=3)["boyutlar"]["Sektör"]
        self.assertEqual((b["hhi"], b["esdeger"], b["en_buyuk"], b["ilk_pay"]), (D("2500.00"), D("4.00"), D("25.00"), D("75.00")))
        r = main.analiz([k("a", 500, "X"), k("b", 300, "y"), k("c", 100, "x"), k("d", 100, "Z"), k("e", 0, "Boş")], ilk=2)
        b = r["boyutlar"]["Sektör"]
        self.assertEqual([(g["ad"], g["kredi"], g["musteri"], g["risk"], g["pay"], g["kumulatif"], g["limit"], g["durum"]) for g in b["satirlar"]],
                         [("X", 2, 2, D(600), D("60.00"), D("60.00"), None, None), ("y", 1, 1, D(300), D("30.00"), D("90.00"), None, None), ("Z", 1, 1, D(100), D("10.00"), D("100.00"), None, None), ("Boş", 1, 1, D(0), D("0.00"), D("100.00"), None, None)])
        self.assertEqual((b["hhi"], b["esdeger"], b["ilk_pay"], r["toplam"], r["uyarilar"]), (D("4600.00"), D("2.17"), D("90.00"), D(1000), []))
        g = main.analiz([k("a", 1, grup="G"), k("b", 2, grup="g"), k("c", 3)])["boyutlar"]["Risk Grubu"]["satirlar"]
        self.assertEqual([(x["ad"], x["kredi"], x["musteri"], x["pay"]) for x in g], [("c", 1, 1, D("50.00")), ("G", 2, 2, D("50.00"))])
        for bos in ([], [k("a", 0)]):
            with self.assertRaises(ValueError):
                main.analiz(bos)

    def test_limitler(self):
        l = [k("a", 300, "A"), k("b", 270, "B"), k("c", "269.99", "C"), k("d", "160.01", "D")]
        r = main.analiz(l, [lim("Sektör", "*", 30), lim("Sektör", "d", 16), lim("Sektör", "Yok", 5), lim("Risk Grubu", "a", 25)])
        self.assertEqual([(g["ad"], g["limit"], g["durum"]) for g in r["boyutlar"]["Sektör"]["satirlar"]], [("A", D(30), "Yaklaştı"), ("B", D(30), "Yaklaştı"), ("C", D(30), "Uygun"), ("D", D(16), "Aşım")])
        self.assertEqual([(g["ad"], g["limit"], g["durum"]) for g in r["boyutlar"]["Risk Grubu"]["satirlar"]], [("a", D(25), "Aşım"), ("b", None, None), ("c", None, None), ("d", None, None)])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         sorted([("Orta", "Yoğunlaşma limitine yaklaştı", "Sektör · A"), ("Orta", "Yoğunlaşma limitine yaklaştı", "Sektör · B"), ("Yüksek", "Yoğunlaşma limiti aşıldı", "Sektör · D"), ("Bilgi", "Limit tanımlı, portföyde yok", "Sektör · Yok"),
                                 ("Yüksek", "Yoğunlaşma limiti aşıldı", "Risk Grubu · a")]))
        self.assertEqual(main.analiz([k("a", 301, "A"), k("b", 699, "B")], [lim("Sektör", "A", 30)])["boyutlar"]["Sektör"]["satirlar"][1]["durum"], "Aşım")

    def test_ozkaynak(self):
        l = [k("a", 250, grup="G1"), k("b", "0.01", grup="G1"), k("c", 250), k("d", 100), k("e", "99.99")]
        r = main.analiz(l, ozkaynak=D(1000))
        self.assertEqual([(g["ad"], g["ozkaynak_oran"]) for g in r["boyutlar"]["Risk Grubu"]["satirlar"]], [("G1", D("25.00")), ("c", D("25.00")), ("d", D("10.00")), ("e", D("10.00"))])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]), [("Bilgi", "Büyük kredi", "Risk Grubu · c"), ("Bilgi", "Büyük kredi", "Risk Grubu · d"), ("Yüksek", "Özkaynak sınırı aşıldı", "Risk Grubu · G1")])
        r = main.analiz(l, ozkaynak=D(1000), ozkaynak_limit=D(20), buyuk_kredi=D(5))
        self.assertEqual(sorted(u["tur"] for u in r["uyarilar"]), ["Büyük kredi"] * 2 + ["Özkaynak sınırı aşıldı"] * 2)
        self.assertEqual(main.analiz(l)["boyutlar"]["Risk Grubu"]["satirlar"][0]["ozkaynak_oran"], None)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "krediler.csv", cikti, ORNEK / "limitler.csv")
            b = r["boyutlar"]
            self.assertEqual((r["toplam"], [(ad, x["adet"], x["hhi"], x["esdeger"], x["en_buyuk"], x["ilk_pay"]) for ad, x in b.items()]),
                             (D(15000000), [("Sektör", 6, D("2208.00"), D("4.53"), D("36.67"), D("92.00")), ("Risk Grubu", 8, D("1719.11"), D("5.82"), D("30.00"), D("82.00")), ("Teminat Türü", 5, D("4664.89"), D("2.14"), D("66.00"), D("100.00"))]))
            self.assertEqual([(g["ad"], g["kredi"], g["risk"], g["pay"], g["limit"], g["durum"]) for g in b["Sektör"]["satirlar"]],
                             [("İnşaat", 3, D(5500000), D("36.67"), D(30), "Aşım"), ("Enerji", 1, D(2500000), D("16.67"), None, None), ("Tekstil", 2, D(2500000), D("16.67"), None, None), ("Turizm", 1, D(1800000), D("12.00"), D(13), "Yaklaştı"),
                              ("Lojistik", 2, D(1500000), D("10.00"), None, None), ("Gıda", 1, D(1200000), D("8.00"), None, None)])
            self.assertEqual([(g["ad"], g["musteri"], g["pay"], g["durum"]) for g in b["Risk Grubu"]["satirlar"]][:3], [("Alfa Grubu", 2, D("30.00"), "Aşım"), ("Eta Enerji", 1, D("16.67"), "Uygun"), ("Beta Tekstil", 1, D("13.33"), "Uygun")])
            self.assertEqual([(g["ad"], g["risk"], g["pay"], g["durum"]) for g in b["Teminat Türü"]["satirlar"]],
                             [("İpotek", D(9900000), D("66.00"), None), ("Çek", D(1700000), D("11.33"), None), ("Kefalet", D(1500000), D("10.00"), None), ("Teminatsız", D(1000000), D("6.67"), "Aşım"), ("Araç Rehni", D(900000), D("6.00"), None)])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Yüksek", "Yoğunlaşma limiti aşıldı", "Risk Grubu · Alfa Grubu"), ("Yüksek", "Yoğunlaşma limiti aşıldı", "Sektör · İnşaat"), ("Yüksek", "Yoğunlaşma limiti aşıldı", "Teminat Türü · Teminatsız"),
                              ("Orta", "Yoğunlaşma limitine yaklaştı", "Sektör · Turizm")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Sektör", "Risk Grubu", "Teminat Türü", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, ly = Path(tmp) / "k.csv", Path(tmp) / "l.csv"
            y.write_text("Unvan;Grup;Faaliyet Alanı;Teminat;Bakiye\nA Ltd.;;;;1.250.000,50\nB A.Ş.;G;İnşaat;İpotek;5\nC;;;;x\nD;;;;-1\n;;;;5\n", encoding="utf-8")
            ly.write_text("Kırılım;Kalem;Limit\nSEKTÖR;İnşaat;12,5\ngrup;;%25\n;;5\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual((l, [(u["tur"], u["kim"]) for u in uy]),
                             ([{"musteri": "A Ltd.", "grup": "A Ltd.", "sektor": "(sektör yok)", "teminat": "(teminat türü yok)", "risk": D("1250000.50")}, {"musteri": "B A.Ş.", "grup": "G", "sektor": "İnşaat", "teminat": "İpotek", "risk": D(5)}],
                              [("Kredi satırı okunamadı", "C"), ("Kredi satırı okunamadı", "D")]))
            self.assertEqual(main.oku_limitler(ly), [{"boyut": "Sektör", "deger": "İnşaat", "limit": D("12.5")}, {"boyut": "Risk Grubu", "deger": "*", "limit": D(25)}])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--krediler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--limitler", str(ly), "--ilk", "10", "--ozkaynak", "5.000.000", "--ozkaynak-limit", "20", "--buyuk-kredi", "5"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--ilk", "0"], ["--ozkaynak", "0"], ["--ozkaynak", "x"], ["--ozkaynak-limit", "5"], ["--buyuk-kredi", "0"], ["--ozkaynak-limit", "101"])], [2] * 6)
            self.assertEqual(main.main(arg + ["--limitler", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Boyut;Değer;Limit\nBölge;A;5\n", "Boyut;Değer;Limit\nSektör;A;0\n", "Boyut;Değer;Limit\nSektör;A;x\n"):
                ly.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg + ["--limitler", str(ly)]), 1)
            for icerik in ("Müşteri;Sektör\nA;B\n", "Müşteri;Risk\n;1\n", "Müşteri;Risk\nA;0\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
