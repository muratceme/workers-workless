"""
Kredi Portföy Yoğunlaşma Analizi — Gerisi Bende kod bloğu
Krediler Tahsis › Kredi Tahsis Müdürü

Kredi portföyünü sektör, risk grubu ve teminat türüne göre dağıtır; yoğunlaşmayı (HHI) ve limit aşımlarını raporlar.
  Pay %                 = boyuttaki değerin riski ÷ toplam risk × 100
  HHI                   = payların (yüzde olarak) karelerinin toplamı; 0 – 10.000 arası, büyüdükçe yoğunlaşma artar
  Eşdeğer sayı          = 10.000 ÷ HHI   (portföy eşit büyüklükte kaç parçaya bölünmüş gibi davranıyor)
  İlk N payı            = en büyük N değerin paylarının toplamı   (--ilk, varsayılan 5)
  Limit aşımı           = pay, limit dosyasındaki yüzdeyi aşıyor;   limite yaklaşma = pay, limitin %90'ı ve üstünde
  Risk grubu ÷ özkaynak = grubun riski ÷ --ozkaynak × 100   (verilirse; --ozkaynak-limit ve --buyuk-kredi ile karşılaştırılır)
Risk grubu boş bırakılan müşteri kendi başına bir grup sayılır. Yasal kredi sınırı hesabındaki ağırlıklar ve istisnalar uygulanmaz.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
K_SUTUN = {"musteri": ("musteri", "musteri adi", "unvan", "firma"), "grup": ("risk grubu", "grup"), "sektor": ("sektor", "faaliyet alani"), "teminat": ("teminat turu", "teminat"), "risk": ("risk", "risk bakiyesi", "bakiye", "kredi riski")}
L_SUTUN = {"boyut": ("boyut", "kirilim"), "deger": ("deger", "ad", "kalem"), "limit": ("limit", "limit %", "limit yuzde")}
BOYUTLAR = (("Sektör", "sektor"), ("Risk Grubu", "grup"), ("Teminat Türü", "teminat"))
BOYUT_AD = {"sektor": "Sektör", "risk grubu": "Risk Grubu", "grup": "Risk Grubu", "teminat turu": "Teminat Türü", "teminat": "Teminat Türü"}
Z, YUZ = Decimal(0), Decimal(100)


def olcu(x) -> Decimal | None:
    """Yüzde okur: virgül ya da tek nokta ondalık ayırıcıdır (12,5 = 12.5)."""
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x or ""))
    try:
        return Decimal(s.replace(".", "").replace(",", ".") if "," in s else s) if s else None
    except InvalidOperation:
        return None


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, K_SUTUN, ("musteri", "risk")):
        ad = metin(r["musteri"])
        if not ad:
            continue
        risk = para(r["risk"])
        if risk is None or risk < 0:
            uy.append(uyari("Yüksek", "Kredi satırı okunamadı", ad, f"Satır {r['_satir']}: risk '{metin(r['risk'])}' okunamadı ya da eksi. Satır alınmadı."))
            continue
        l.append({"musteri": ad, "grup": metin(r["grup"]) or ad, "sektor": metin(r["sektor"]) or "(sektör yok)", "teminat": metin(r["teminat"]) or "(teminat türü yok)", "risk": risk})
    return l, uy


def oku_limitler(yol: Path) -> list[dict]:
    l = []
    for r in kayitlar(yol, L_SUTUN, ("boyut", "limit")):
        if not metin(r["boyut"]):
            continue
        boyut, limit = BOYUT_AD.get(katla(r["boyut"])), olcu(r["limit"])
        if boyut is None or limit is None or not 0 < limit <= 100:
            raise ValueError(f"{yol.name} satır {r['_satir']}: boyut Sektör, Risk Grubu ya da Teminat Türü; limit 0–100 arasında bir yüzde olmalı.")
        l.append({"boyut": boyut, "deger": metin(r["deger"]) or "*", "limit": limit})
    return l


def analiz(krediler: list[dict], limitler: list[dict] | None = None, ilk: int = 5, ozkaynak: Decimal | None = None, ozkaynak_limit: Decimal = Decimal(25), buyuk_kredi: Decimal = Decimal(10)) -> dict:
    toplam = sum((k["risk"] for k in krediler), Z)
    if not toplam:
        raise ValueError("Analiz edilecek risk yok (toplam risk sıfır).")
    uy, boyutlar = [], {}
    for ad, alan in BOYUTLAR:
        grup = {}
        for k in krediler:
            g = grup.setdefault(katla(k[alan]), {"ad": k[alan], "kredi": 0, "risk": Z, "musteriler": set()})
            g["kredi"] += 1
            g["risk"] += k["risk"]
            g["musteriler"].add(katla(k["musteri"]))
        satirlar = sorted(grup.values(), key=lambda g: (-g["risk"], katla(g["ad"])))
        genel = next((x["limit"] for x in limitler or [] if x["boyut"] == ad and x["deger"] == "*"), None)
        ozel = {katla(x["deger"]): x["limit"] for x in limitler or [] if x["boyut"] == ad and x["deger"] != "*"}
        kum, hhi = Z, Z
        for g in satirlar:
            ham = g["risk"] * YUZ / toplam
            kum += ham
            hhi += ham * ham
            g["musteri"], g["pay"], g["kumulatif"] = len(g.pop("musteriler")), yuvarla(ham), yuvarla(kum)
            g["limit"] = ozel.get(katla(g["ad"]), genel)
            g["durum"] = None if g["limit"] is None else "Aşım" if ham > g["limit"] else "Yaklaştı" if ham >= g["limit"] * Decimal("0.9") else "Uygun"
            kim = f"{ad} · {g['ad']}"
            if g["durum"] == "Aşım":
                uy.append(uyari("Yüksek", "Yoğunlaşma limiti aşıldı", kim, f"Pay %{g['pay']}, limit %{g['limit']:f}. Risk {g['risk']:,.2f}; limite inmek için {g['risk'] - toplam * g['limit'] / YUZ:,.2f} azalmalı (toplam sabitken)."))
            elif g["durum"] == "Yaklaştı":
                uy.append(uyari("Orta", "Yoğunlaşma limitine yaklaştı", kim, f"Pay %{g['pay']}, limit %{g['limit']:f}. Kalan alan {toplam * g['limit'] / YUZ - g['risk']:,.2f}."))
            if ad == "Risk Grubu" and ozkaynak:
                g["ozkaynak_oran"] = yuvarla(g["risk"] * YUZ / ozkaynak)
                if g["risk"] * YUZ / ozkaynak > ozkaynak_limit:
                    uy.append(uyari("Yüksek", "Özkaynak sınırı aşıldı", kim, f"Grup riski {g['risk']:,.2f}, özkaynağın %{g['ozkaynak_oran']}'i (sınır %{ozkaynak_limit:f})."))
                elif g["risk"] * YUZ / ozkaynak >= buyuk_kredi:
                    uy.append(uyari("Bilgi", "Büyük kredi", kim, f"Grup riski {g['risk']:,.2f}, özkaynağın %{g['ozkaynak_oran']}'i (büyük kredi eşiği %{buyuk_kredi:f})."))
            else:
                g["ozkaynak_oran"] = None
        boyutlar[ad] = {"satirlar": satirlar, "adet": len(satirlar), "hhi": yuvarla(hhi), "esdeger": yuvarla(Decimal(10000) / hhi), "en_buyuk": satirlar[0]["pay"], "ilk_pay": yuvarla(sum((g["risk"] for g in satirlar[:ilk]), Z) * YUZ / toplam)}
        for x in limitler or []:
            if x["boyut"] == ad and x["deger"] != "*" and katla(x["deger"]) not in grup:
                uy.append(uyari("Bilgi", "Limit tanımlı, portföyde yok", f"{ad} · {x['deger']}", f"Limit dosyasında %{x['limit']:f} limit var ama portföyde bu değerde risk yok (yazım farkı olabilir)."))
    return {"krediler": krediler, "toplam": toplam, "boyutlar": boyutlar, "ilk": ilk, "ozkaynak": ozkaynak, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    _baslik(ws, ["Boyut", "Değer Sayısı", "HHI", "Eşdeğer Sayı", "En Büyük Pay %", f"İlk {r['ilk']} Payı %"], (20, 11, 12, 12, 13, 13))
    for ad, b in r["boyutlar"].items():
        _satir(ws, [ad, b["adet"], b["hhi"], b["esdeger"], b["en_buyuk"], b["ilk_pay"]], para_s=(3, 4, 5, 6))
    ws.append([])
    ws.append([f"Toplam risk: {r['toplam']:,.2f} · {len(r['krediler'])} kredi" + (f" · özkaynak {r['ozkaynak']:,.2f}" if r["ozkaynak"] else "")])
    ws.append(["HHI = payların (yüzde) karelerinin toplamı; 10.000 tek bir değerde tam yoğunlaşma demektir. Eşdeğer sayı = 10.000 ÷ HHI."])
    for ad, b in r["boyutlar"].items():
        ws = wb.create_sheet(ad)
        _baslik(ws, [ad, "Kredi", "Müşteri", "Risk", "Pay %", "Kümülatif %", "Limit %", "Durum"] + (["Özkaynağa Oran %"] if ad == "Risk Grubu" and r["ozkaynak"] else []), (30, 8, 9, 18, 9, 11, 9, 11, 12))
        for g in b["satirlar"]:
            _satir(ws, [g["ad"], g["kredi"], g["musteri"], g["risk"], g["pay"], g["kumulatif"], g["limit"], g["durum"]] + ([g["ozkaynak_oran"]] if ad == "Risk Grubu" and r["ozkaynak"] else []), para_s=(4, 5, 6, 7, 9),
                   dolgu={8: {"Aşım": KIRMIZI, "Yaklaştı": SARI, "Uygun": YESIL}.get(g["durum"])})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, limit_yol: Path | None = None, **kw) -> dict:
    krediler, uy = oku(yol)
    r = analiz(krediler, oku_limitler(limit_yol) if limit_yol else None, **kw)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Kredi portföyünün sektör, risk grubu ve teminat türü yoğunlaşmasını (HHI) ve limit aşımlarını raporlar.")
    p.add_argument("--krediler", type=Path, help="Müşteri, Risk Grubu, Sektör, Teminat Türü, Risk")
    p.add_argument("--limitler", type=Path, help="Boyut (Sektör / Risk Grubu / Teminat Türü), Değer (* = boyuttaki her değer), Limit yüzdesi (isteğe bağlı)")
    p.add_argument("--ilk", type=int, default=5, help="En büyük kaç değerin toplam payı raporlansın (varsayılan 5)")
    p.add_argument("--ozkaynak", help="Özkaynak tutarı; verilirse risk gruplarının özkaynağa oranı hesaplanır")
    p.add_argument("--ozkaynak-limit", default="25", help="Bir risk grubunun özkaynağa oranı için üst sınır, yüzde (varsayılan 25)")
    p.add_argument("--buyuk-kredi", default="10", help="Büyük kredi eşiği, özkaynağın yüzdesi (varsayılan 10)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kredi_portfoy_yogunlasma.xlsx")
    a = p.parse_args(argv)
    ornek = a.krediler is None
    yol = ORNEK / "krediler.csv" if ornek else a.krediler
    limit_yol = ORNEK / "limitler.csv" if ornek and a.limitler is None else a.limitler
    ozkaynak, ol, bk = para(a.ozkaynak), olcu(a.ozkaynak_limit), olcu(a.buyuk_kredi)
    if a.ilk < 1 or (a.ozkaynak and (ozkaynak is None or ozkaynak <= 0)) or ol is None or bk is None or not 0 < bk <= ol <= 100:
        print("[X] --ilk en az 1; --ozkaynak pozitif; 0 < --buyuk-kredi ≤ --ozkaynak-limit ≤ 100 olmalı")
        return 2
    for y in (yol, limit_yol):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(yol, a.cikti, limit_yol, ilk=a.ilk, ozkaynak=ozkaynak, ozkaynak_limit=ol, buyuk_kredi=bk)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(r['krediler'])} kredi · toplam risk {r['toplam']:,.2f} · " + " · ".join(f"{ad} HHI {b['hhi']} (en büyük %{b['en_buyuk']})" for ad, b in r["boyutlar"].items()))
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
