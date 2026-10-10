"""
Proforma Fatura Hazırlama — Workers / Workless kod bloğu
Dış Ticaret › Dış Ticaret Uzmanı

Sipariş, fiyat ve teslim şekli (Incoterms 2020) bilgilerinden proforma fatura ve çeki listesi hazırlar:
  - Satırlar: ürün kartından İngilizce tanım, GTİP, menşe, birim fiyat (siparişte fiyat yazılıysa o); tutar = miktar ×
    birim fiyat. Ürün dövizi proforma dövizinden farklıysa satır alınmaz (çevrim yapılmaz).
  - Teslim şekli kontrolü: Incoterms 2020 kuralları (EXW, FCA, CPT, CIP, DAP, DPU, DDP her taşıma türü; FAS, FOB, CFR,
    CIF yalnız deniz ve iç su yolu). Teslim yeri zorunlu. Navlun C ve D gruplarında, sigorta CIF ve CIP'te satıcıya
    aittir: eksikse uyarı, gereksizse (ör. FOB'da navlun) uyarı. DAT yazılırsa DPU önerilir.
  - Toplam: mal bedeli + navlun + sigorta (teslim şekline göre) ve tutarın İngilizce yazıyla karşılığı.
  - Çeki listesi: koli sayısı = ⌈miktar / koli içi adet⌉, net ağırlık = miktar × birim net ağırlık, brüt = net + koli ×
    koli darası, hacim = koli × ölçüler (m³); koli numaraları sıralı.
  - Kontroller: GTİP 6–12 hane, menşe, ağırlık / koli bilgisi, banka bilgileri, geçerlilik.
Çıktı: müşteriye gönderilecek dosya (Proforma Invoice + Packing List, İngilizce) ve ayrı iç kontrol dosyası.
İnternete bağlanmaz.

Kullanım:
    python main.py                                                  # örnek: 6 kalem, CIF Hamburg
    python main.py --siparis siparis.xlsx --urunler urunler.xlsx --bilgi proforma_bilgisi.csv --cikti proformalar
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
K2 = Decimal("0.01")
SIFIR = Decimal(0)

# Incoterms 2020: kod → (ad, yalnız deniz/iç su yolu, navlun satıcıda, sigorta satıcıda)
INCOTERMS = {
    "EXW": ("Ex Works", False, False, False), "FCA": ("Free Carrier", False, False, False), "CPT": ("Carriage Paid To", False, True, False),
    "CIP": ("Carriage and Insurance Paid To", False, True, True), "DAP": ("Delivered at Place", False, True, False),
    "DPU": ("Delivered at Place Unloaded", False, True, False), "DDP": ("Delivered Duty Paid", False, True, False),
    "FAS": ("Free Alongside Ship", True, False, False), "FOB": ("Free On Board", True, False, False), "CFR": ("Cost and Freight", True, True, False),
    "CIF": ("Cost, Insurance and Freight", True, True, True),
}
DENIZ = {"deniz", "denizyolu", "deniz yolu", "sea", "gemi", "ic su yolu"}

URUN_SUTUNLARI = {"kod": ("urun kodu", "kod", "stok kodu"), "tanim": ("tanim", "ingilizce tanim", "description", "urun"), "gtip": ("gtip", "hs code", "hs kodu"),
                  "mense": ("mense", "origin", "mense ulke"), "birim": ("birim", "unit"), "fiyat": ("birim fiyat", "fiyat", "unit price"),
                  "doviz": ("doviz", "para birimi", "currency"), "net": ("net agirlik", "birim net agirlik", "net agirlik kg", "birim net agirlik kg"),
                  "koli_ici": ("koli ici adet", "koli ici", "koli adedi"), "dara": ("koli darasi", "koli darasi kg", "koli dara", "dara"),
                  "olcu": ("koli olculeri", "koli olculeri cm", "koli olcusu", "olculer")}
SIPARIS_SUTUNLARI = {"kod": ("urun kodu", "kod", "stok kodu"), "miktar": ("miktar", "adet", "quantity"), "fiyat": ("birim fiyat", "fiyat", "anlasilan fiyat"),
                     "not": ("not", "aciklama")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def en_tutar(x: Decimal) -> str:
    return f"{x.quantize(K2, ROUND_HALF_UP):,.2f}"


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


# ----------------------------------------------------------------------------
# Tutarın İngilizce yazıyla karşılığı
# ----------------------------------------------------------------------------

_BIRLER = ["", "ONE", "TWO", "THREE", "FOUR", "FIVE", "SIX", "SEVEN", "EIGHT", "NINE", "TEN", "ELEVEN", "TWELVE", "THIRTEEN", "FOURTEEN", "FIFTEEN", "SIXTEEN",
           "SEVENTEEN", "EIGHTEEN", "NINETEEN"]
_ONLAR = ["", "", "TWENTY", "THIRTY", "FORTY", "FIFTY", "SIXTY", "SEVENTY", "EIGHTY", "NINETY"]


def _uc_hane(n: int) -> str:
    p = []
    if n >= 100:
        p.append(f"{_BIRLER[n // 100]} HUNDRED")
        n %= 100
    if n >= 20:
        p.append(_ONLAR[n // 10] + (f"-{_BIRLER[n % 10]}" if n % 10 else ""))
    elif n:
        p.append(_BIRLER[n])
    return " ".join(p)


def yaziyla(tutar: Decimal, doviz: str) -> str:
    """1234.50 USD → 'SAY USD ONE THOUSAND TWO HUNDRED THIRTY-FOUR AND 50/100 ONLY'."""
    t = tutar.quantize(K2, ROUND_HALF_UP)
    tam, kurus = int(t), int((t - int(t)) * 100)
    if tam == 0:
        kelime = "ZERO"
    else:
        parca, i = [], 0
        olcek = ["", " THOUSAND", " MILLION", " BILLION"]
        while tam:
            tam, kalan = divmod(tam, 1000)
            if kalan:
                parca.append(_uc_hane(kalan) + olcek[i])
            i += 1
        kelime = " ".join(reversed(parca))
    return f"SAY {doviz} {kelime} AND {kurus:02d}/100 ONLY"


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

def bilgi_oku(yol: Path) -> dict:
    s = {}
    for r in tablo_oku(yol):
        if len(r) >= 2 and metin(r[0]):
            s[katla(r[0])] = r[1]
    al = lambda *adlar: next((metin(s[a]) for a in adlar if metin(s.get(a))), "")  # noqa: E731
    return {"no": al("proforma no", "no") or "PROFORMA", "tarih": tarih(s.get("tarih")) or date.today(), "gecerlilik": int(para(s.get("gecerlilik gun")) or 30),
            "satici": al("satici"), "satici_adres": al("satici adresi"), "alici": al("alici"), "alici_adres": al("alici adresi"), "ulke": al("ulke", "alici ulkesi"),
            "incoterm": al("teslim sekli", "incoterms").upper().split()[0] if al("teslim sekli", "incoterms") else "", "teslim_yeri": al("teslim yeri"),
            "tasima": al("tasima sekli"), "odeme": al("odeme sekli", "odeme"), "doviz": (al("doviz") or "USD").upper(), "navlun": para(s.get("navlun")),
            "sigorta": para(s.get("sigorta")), "banka": al("banka"), "iban": al("iban"), "swift": al("swift"), "yukleme": al("yukleme limani", "yukleme yeri"),
            "varis": al("varis limani", "varis yeri"), "sevk": al("tahmini sevk", "sevk tarihi"), "notlar": al("notlar")}


@dataclass
class Urun:
    kod: str
    tanim: str
    gtip: str
    mense: str
    birim: str
    fiyat: Decimal | None
    doviz: str
    net: Decimal | None
    koli_ici: Decimal | None
    dara: Decimal | None
    olcu: tuple[Decimal, Decimal, Decimal] | None


@dataclass
class Kalem:
    sira: int
    urun: Urun
    miktar: Decimal
    fiyat: Decimal
    tutar: Decimal = SIFIR
    koli: int | None = None
    net: Decimal | None = None
    brut: Decimal | None = None
    hacim: Decimal | None = None
    koli_no: str = ""
    notlar: list[str] = field(default_factory=list)


def urun_oku(yol: Path) -> dict[str, Urun]:
    u = {}
    for r in kayitlar(yol, URUN_SUTUNLARI, ("kod",)):
        kod = metin(r.get("kod")).upper()
        if not kod:
            continue
        olcu = [Decimal(v.replace(",", ".")) for v in re.findall(r"\d+(?:[.,]\d+)?", metin(r.get("olcu")))]
        u[kod] = Urun(kod, metin(r.get("tanim")), re.sub(r"\D", "", metin(r.get("gtip"))), metin(r.get("mense")), metin(r.get("birim")) or "PCS", para(r.get("fiyat")),
                      (metin(r.get("doviz")) or "").upper(), para(r.get("net")), para(r.get("koli_ici")), para(r.get("dara")),
                      tuple(olcu[:3]) if len(olcu) >= 3 else None)
    return u


# ----------------------------------------------------------------------------
# Hesap ve kontroller
# ----------------------------------------------------------------------------

def hazirla(urunler: dict[str, Urun], siparis: list[dict], b: dict) -> dict:
    uy, kalemler = [], []

    def uyar(onem, tur, kim, aciklama):
        uy.append({"onem": onem, "tur": tur, "kim": kim, "aciklama": aciklama})

    for r in siparis:
        kod, m = metin(r.get("kod")).upper(), para(r.get("miktar"))
        if not kod:
            continue
        u = urunler.get(kod)
        if not u:
            uyar("Yüksek", "Ürün kartı yok", kod, "Kalem proformaya alınmadı")
            continue
        if not m or m <= 0:
            uyar("Yüksek", "Geçersiz miktar", kod, f"Miktar '{metin(r.get('miktar'))}'; kalem alınmadı")
            continue
        fiyat = para(r.get("fiyat")) if para(r.get("fiyat")) is not None else u.fiyat
        if fiyat is None:
            uyar("Yüksek", "Fiyat yok", kod, "Siparişte ve ürün kartında birim fiyat yok; kalem alınmadı")
            continue
        if u.doviz and u.doviz != b["doviz"] and para(r.get("fiyat")) is None:
            uyar("Yüksek", "Döviz uyuşmuyor", kod, f"Ürün fiyatı {u.doviz}, proforma {b['doviz']}; çevrim yapılmaz. Siparişe {b['doviz']} fiyat yazın")
            continue
        k = Kalem(len(kalemler) + 1, u, m, fiyat)
        k.tutar = (m * fiyat).quantize(K2, ROUND_HALF_UP)
        if para(r.get("fiyat")) is not None and u.fiyat is not None and u.doviz == b["doviz"] and fiyat < u.fiyat:
            uyar("Bilgi", "Liste altı fiyat", kod, f"Sipariş fiyatı {en_tutar(fiyat)}, liste {en_tutar(u.fiyat)} {b['doviz']}")
        if not 6 <= len(u.gtip) <= 12:
            uyar("Orta", "GTİP eksik / hatalı", kod, f"GTİP '{u.gtip or '—'}'; 6 (HS) ile 12 hane (GTİP) arasında olmalı")
        if not u.mense:
            uyar("Orta", "Menşe yok", kod, "Menşe ülke yazılmamış")
        if not u.tanim:
            uyar("Orta", "Tanım yok", kod, "İngilizce ürün tanımı yok")
        if u.net is not None:
            k.net = m * u.net
        else:
            uyar("Orta", "Ağırlık yok", kod, "Birim net ağırlık yok; çeki listesinde ağırlık boş kalır")
        if u.koli_ici:
            k.koli = math.ceil(m / u.koli_ici)
            if m % u.koli_ici:
                k.notlar.append(f"last carton {m % u.koli_ici:g} {u.birim}")
                uyar("Bilgi", "Kısmi koli", kod, f"{m:g} / {u.koli_ici:g}: son koli tam değil ({m % u.koli_ici:g} {u.birim})")
            if k.net is not None and u.dara is not None:
                k.brut = k.net + k.koli * u.dara
            if u.olcu:
                k.hacim = (Decimal(k.koli) * u.olcu[0] * u.olcu[1] * u.olcu[2] / Decimal(1000000)).quantize(Decimal("0.001"), ROUND_HALF_UP)
        else:
            uyar("Orta", "Koli bilgisi yok", kod, "Koli içi adet yok; koli sayısı ve brüt ağırlık hesaplanamadı")
        kalemler.append(k)
    if not kalemler:
        raise ValueError("Proformaya alınacak kalem yok")
    no = 1
    for k in kalemler:
        if k.koli:
            k.koli_no = f"{no}–{no + k.koli - 1}" if k.koli > 1 else str(no)
            no += k.koli

    # Teslim şekli
    kod = b["incoterm"]
    kural = INCOTERMS.get(kod)
    deniz = katla(b["tasima"]) in DENIZ
    navlun = b["navlun"] or SIFIR
    sigorta = b["sigorta"] or SIFIR
    if not kod:
        uyar("Yüksek", "Teslim şekli yok", "Incoterms", "Teslim şekli yazılmamış")
    elif kod == "DAT":
        uyar("Yüksek", "Eski teslim şekli", "DAT", "DAT, Incoterms 2020'de DPU (Delivered at Place Unloaded) olarak değişti")
    elif not kural:
        uyar("Yüksek", "Bilinmeyen teslim şekli", kod, "Incoterms 2020: EXW, FCA, FAS, FOB, CFR, CIF, CPT, CIP, DAP, DPU, DDP")
    else:
        if kural[1] and b["tasima"] and not deniz:
            alternatif = {"FOB": "FCA", "FAS": "FCA", "CFR": "CPT", "CIF": "CIP"}[kod]
            uyar("Yüksek", "Teslim şekli taşıma türüne uymuyor", kod,
                 f"{kod} yalnız deniz ve iç su yolu taşımasında kullanılır; taşıma '{b['tasima']}'. {alternatif} kullanmayı değerlendirin")
        if kural[1] and deniz and kod in ("FOB", "CFR", "CIF") and "konteyner" in katla(b["notlar"] + " " + b["tasima"]):
            uyar("Bilgi", "Konteyner yüklemesi", kod, "Konteynerli yüklemede mal terminalde teslim edildiğinden FCA / CPT / CIP önerilir")
        if not b["teslim_yeri"]:
            uyar("Yüksek", "Teslim yeri yok", kod, "Incoterms kuralı belirli bir yer / liman ile yazılır (ör. 'CIF Hamburg Port, Incoterms 2020')")
        if kural[2] and not navlun:
            uyar("Orta", "Navlun yok", kod, f"{kod} teslimde navlun satıcıya aittir; navlun tutarı girilmemiş (fiyata dahilse notlara yazın)")
        if not kural[2] and navlun:
            uyar("Orta", "Gereksiz navlun", kod, f"{kod} teslimde ana taşıma alıcıya aittir; navlun satırı proformaya eklenmedi")
            navlun = SIFIR
        if kural[3] and not sigorta:
            uyar("Orta", "Sigorta yok", kod, f"{kod} teslimde satıcı sigorta yaptırır (en az mal bedelinin %110'u; CIF için ICC (C), CIP için ICC (A)); tutar girilmemiş")
        if not kural[3] and sigorta:
            uyar("Bilgi", "Sigorta satırı", kod, f"{kod} teslimde satıcının sigorta yükümlülüğü yoktur; sigorta satırı proformaya eklenmedi")
            sigorta = SIFIR
    if not b["alici"]:
        uyar("Yüksek", "Alıcı yok", "Bilgi", "Alıcı unvanı yazılmamış")
    if not b["odeme"]:
        uyar("Orta", "Ödeme şekli yok", "Bilgi", "Ödeme şekli yazılmamış (ör. peşin, mal mukabili, vesaik mukabili, akreditif)")
    if not (b["banka"] and b["iban"] and b["swift"]):
        uyar("Orta", "Banka bilgisi eksik", "Bilgi", "Banka, IBAN ve SWIFT bilgilerini tamamlayın")
    mal = sum((k.tutar for k in kalemler), SIFIR)
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    uy.sort(key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))
    toplam = mal + navlun + sigorta
    return {"kalemler": kalemler, "bilgi": b, "mal": mal, "navlun": navlun, "sigorta": sigorta, "toplam": toplam, "yaziyla": yaziyla(toplam, b["doviz"]),
            "uyarilar": uy, "gecerlilik_tarihi": b["tarih"] + timedelta(days=b["gecerlilik"]),
            "koli": sum(k.koli or 0 for k in kalemler), "net": sum((k.net for k in kalemler if k.net is not None), SIFIR),
            "brut": sum((k.brut for k in kalemler if k.brut is not None), SIFIR), "hacim": sum((k.hacim for k in kalemler if k.hacim is not None), SIFIR),
            "teslim": f"{kod} {b['teslim_yeri']}, Incoterms® 2020".strip() if kod else ""}


# ----------------------------------------------------------------------------
# Çıktılar
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
INCE = Side(style="thin", color="BBBBBB")
CERCEVE = Border(top=INCE, bottom=INCE, left=INCE, right=INCE)
PF = "#,##0.00"


def _ust_bilgi(ws, s, baslik):
    b = s["bilgi"]
    ws.append([b["satici"]])
    ws["A1"].font = Font(bold=True, size=13)
    if b["satici_adres"]:
        ws.append([b["satici_adres"]])
    ws.append([])
    ws.append([baslik])
    ws.cell(ws.max_row, 1).font = Font(bold=True, size=14)
    ws.append([])
    for etiket, deger in (("No", b["no"]), ("Date", b["tarih"].strftime("%d.%m.%Y")), ("Buyer", b["alici"]), ("Address", b["alici_adres"]), ("Country", b["ulke"]),
                          ("Delivery Terms", s["teslim"]), ("Mode of Transport", b["tasima"]), ("Port / Place of Loading", b["yukleme"]),
                          ("Port / Place of Discharge", b["varis"]), ("Estimated Shipment", b["sevk"])):
        if deger:
            ws.append([etiket, None, deger])
            ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.append([])


def _tablo_basligi(ws, basliklar):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font, h.border = BASLIK_DOLGU, BASLIK_YAZI, CERCEVE
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def proforma_yaz(yol: Path, s: dict) -> None:
    b, dv = s["bilgi"], s["bilgi"]["doviz"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Proforma Invoice"
    for j, w in enumerate((6, 12, 44, 14, 10, 10, 7, 13, 15), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    _ust_bilgi(ws, s, "PROFORMA INVOICE")
    _tablo_basligi(ws, ["No", "Item Code", "Description of Goods", "HS Code", "Origin", "Quantity", "Unit", f"Unit Price ({dv})", f"Amount ({dv})"])
    for k in s["kalemler"]:
        ws.append([k.sira, k.urun.kod, k.urun.tanim, k.urun.gtip, k.urun.mense, float(k.miktar), k.urun.birim, float(k.fiyat), float(k.tutar)])
        for j in range(1, 10):
            ws.cell(ws.max_row, j).border = CERCEVE
        ws.cell(ws.max_row, 8).number_format = ws.cell(ws.max_row, 9).number_format = PF
        ws.cell(ws.max_row, 3).alignment = Alignment(wrap_text=True)
    ws.append([])
    satirlar = [("Total value of goods", s["mal"])]
    if s["navlun"]:
        satirlar.append(("Freight", s["navlun"]))
    if s["sigorta"]:
        satirlar.append(("Insurance", s["sigorta"]))
    satirlar.append((f"TOTAL {b['incoterm']} {b['teslim_yeri']} ({dv})".replace("  ", " "), s["toplam"]))
    for etiket, v in satirlar:
        ws.append([None, None, None, None, None, etiket, None, None, float(v)])
        ws.cell(ws.max_row, 9).number_format = PF
        if etiket.startswith("TOTAL"):
            ws.cell(ws.max_row, 6).font = ws.cell(ws.max_row, 9).font = Font(bold=True)
    ws.append([])
    ws.append([s["yaziyla"]])
    ws.cell(ws.max_row, 1).font = Font(italic=True)
    ws.append([])
    for etiket, deger in (("Payment Terms", b["odeme"]), ("Validity", f"This proforma invoice is valid until {s['gecerlilik_tarihi']:%d.%m.%Y}."),
                          ("Packing", f"{s['koli']} cartons; net {s['net']:,.2f} kg; gross {s['brut']:,.2f} kg; {s['hacim']:,.3f} m³" if s["koli"] else ""),
                          ("Bank", b["banka"]), ("IBAN", b["iban"]), ("SWIFT", b["swift"]), ("Remarks", b["notlar"])):
        if deger:
            ws.append([etiket, None, deger])
            ws.cell(ws.max_row, 1).font = Font(bold=True)
            ws.cell(ws.max_row, 3).alignment = Alignment(wrap_text=True)
    ws.append([])
    ws.append(["This is a proforma invoice and not a demand for payment unless otherwise agreed. It is not a commercial invoice."])
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    pl = wb.create_sheet("Packing List")
    for j, w in enumerate((6, 12, 44, 10, 7, 9, 12, 12, 12, 11, 16), 1):
        pl.column_dimensions[get_column_letter(j)].width = w
    _ust_bilgi(pl, s, "PACKING LIST")
    _tablo_basligi(pl, ["No", "Item Code", "Description of Goods", "Quantity", "Unit", "Cartons", "Carton No", "Net Weight (kg)", "Gross Weight (kg)",
                        "Volume (m³)", "Remarks"])
    for k in s["kalemler"]:
        pl.append([k.sira, k.urun.kod, k.urun.tanim, float(k.miktar), k.urun.birim, k.koli, k.koli_no, None if k.net is None else float(k.net),
                   None if k.brut is None else float(k.brut), None if k.hacim is None else float(k.hacim), "; ".join(k.notlar)])
        for j in range(1, 12):
            pl.cell(pl.max_row, j).border = CERCEVE
        pl.cell(pl.max_row, 8).number_format = pl.cell(pl.max_row, 9).number_format = PF
        pl.cell(pl.max_row, 10).number_format = "0.000"
    pl.append([None, None, "TOTAL", None, None, s["koli"], None, float(s["net"]), float(s["brut"]), float(s["hacim"])])
    for j in (3, 6, 8, 9, 10):
        pl.cell(pl.max_row, j).font = Font(bold=True)
    pl.cell(pl.max_row, 8).number_format = pl.cell(pl.max_row, 9).number_format = PF
    pl.cell(pl.max_row, 10).number_format = "0.000"
    pl.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(yol)


def kontrol_yaz(yol: Path, s: dict) -> None:
    b = s["bilgi"]
    wb = Workbook()
    oz = wb.active
    oz.title = "Özet"
    oz.column_dimensions["A"].width, oz.column_dimensions["B"].width = 30, 60
    yuksek = sum(u["onem"] == "Yüksek" for u in s["uyarilar"])
    for e, v in [("Proforma No", b["no"]), ("Alıcı", f"{b['alici']} ({b['ulke']})"), ("Teslim şekli", s["teslim"]), ("Taşıma", b["tasima"]),
                 ("Mal bedeli", f"{en_tutar(s['mal'])} {b['doviz']}"), ("Navlun", f"{en_tutar(s['navlun'])} {b['doviz']}"),
                 ("Sigorta", f"{en_tutar(s['sigorta'])} {b['doviz']}"), ("Toplam", f"{en_tutar(s['toplam'])} {b['doviz']}"),
                 ("Koli / net / brüt / hacim", f"{s['koli']} koli · {s['net']:.2f} kg · {s['brut']:.2f} kg · {s['hacim']:.3f} m³"),
                 ("Gönderime hazır mı?", "HAYIR — yüksek önemli uyarıları giderin" if yuksek else "Evet (uyarıları gözden geçirin)")]:
        oz.append([e, v])
    oz.cell(oz.max_row, 2).fill = PatternFill("solid", fgColor="FDE2E1" if yuksek else "E3F4E1")
    oz.append([])
    oz.append(["Bu dosya şirket içi kontrol içindir; alıcıya gönderilmez."])
    uy = wb.create_sheet("Uyarılar")
    uy.append(["Önem", "Tür", "Kalem / Alan", "Açıklama"])
    for h in uy[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for j, w in enumerate((9, 30, 16, 110), 1):
        uy.column_dimensions[get_column_letter(j)].width = w
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["kim"], u["aciklama"]])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(yol)


def calistir(siparis_yolu: Path, urun_yolu: Path, bilgi_yolu: Path, cikti_klasoru: Path) -> dict:
    s = hazirla(urun_oku(urun_yolu), kayitlar(siparis_yolu, SIPARIS_SUTUNLARI, ("kod", "miktar")), bilgi_oku(bilgi_yolu))
    cikti_klasoru.mkdir(parents=True, exist_ok=True)
    ad = re.sub(r"[^a-z0-9]+", "-", katla(s["bilgi"]["no"])).strip("-") or "proforma"
    s["proforma"] = cikti_klasoru / f"proforma_{ad}.xlsx"
    s["kontrol"] = cikti_klasoru / f"proforma_{ad}_kontrol.xlsx"
    proforma_yaz(s["proforma"], s)
    kontrol_yaz(s["kontrol"], s)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Sipariş ve teslim şekli bilgilerinden proforma fatura ve çeki listesi hazırlar.")
    p.add_argument("--siparis", type=Path, default=ORNEK / "siparis.csv", help="Ürün Kodu, Miktar, Birim Fiyat (anlaşılan; boşsa ürün kartındaki), Not")
    p.add_argument("--urunler", type=Path, default=ORNEK / "urunler.csv",
                   help="Ürün Kodu, Tanım (İngilizce), GTİP, Menşe, Birim, Birim Fiyat, Döviz, Net Ağırlık (kg), Koli İçi Adet, Koli Darası (kg), "
                        "Koli Ölçüleri (cm, 60x40x35)")
    p.add_argument("--bilgi", type=Path, default=ORNEK / "proforma_bilgisi.csv",
                   help="Alan;Değer: Proforma No, Tarih, Satıcı, Alıcı, Ülke, Teslim Şekli, Teslim Yeri, Taşıma Şekli, Ödeme Şekli, Döviz, Navlun, Sigorta, Banka...")
    p.add_argument("--cikti", type=Path, default=Path("cikti"), help="Çıktı klasörü")
    a = p.parse_args(argv)
    for y in (a.siparis, a.urunler, a.bilgi):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        s = calistir(a.siparis, a.urunler, a.bilgi, a.cikti)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    b = s["bilgi"]
    print(f"[OK] {b['no']} · {b['alici']} · {len(s['kalemler'])} kalem · {s['teslim']} · toplam {en_tutar(s['toplam'])} {b['doviz']}")
    print(f"[OK] Çeki listesi: {s['koli']} koli · net {s['net']:.2f} kg · brüt {s['brut']:.2f} kg · {s['hacim']:.3f} m³")
    for u in s["uyarilar"]:
        if u["onem"] != "Bilgi":
            print(f"[!] {u['kim']} · {u['tur']}: {u['aciklama']}")
    print(f"[OK] Proforma: {s['proforma'].resolve()}")
    print(f"[OK] İç kontrol: {s['kontrol'].resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
