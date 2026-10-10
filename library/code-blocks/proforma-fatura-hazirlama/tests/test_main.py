import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def bilgi(tmp, **degisen):
    satirlar = {}
    for r in main.tablo_oku(ORNEK / "proforma_bilgisi.csv")[1:]:
        satirlar[r[0]] = r[1] if len(r) > 1 else ""
    satirlar.update(degisen)
    yol = Path(tmp) / "b.csv"
    yol.write_text("Alan;Değer\n" + "\n".join(f"{k};{v}" for k, v in satirlar.items()) + "\n", encoding="utf-8")
    return yol


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.s = main.calistir(ORNEK / "siparis.csv", ORNEK / "urunler.csv", ORNEK / "proforma_bilgisi.csv", Path(cls.tmp.name))
        cls.u = {(u["tur"], u["kim"]) for u in cls.s["uyarilar"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_tutarlar(self):
        s = self.s
        self.assertEqual(s["mal"], D("82662.00"))               # 15.600 + 14.880 + 9.750 + 17.920 + 11.312 + 13.200
        self.assertEqual((s["navlun"], s["sigorta"], s["toplam"]), (D("2850.00"), D(0), D("85512.00")))
        self.assertEqual(s["yaziyla"], "SAY EUR EIGHTY-FIVE THOUSAND FIVE HUNDRED TWELVE AND 00/100 ONLY")
        self.assertEqual(s["kalemler"][0].fiyat, D("6.50"))     # siparişteki anlaşılan fiyat

    def test_ceki_listesi(self):
        s = self.s
        self.assertEqual((s["koli"], s["net"], s["brut"], s["hacim"]), (431, D("5983.500"), D("6539.800"), D("46.056")))
        k = s["kalemler"][4]                                    # 1010 / 20 = 51 koli, son koli 10 adet
        self.assertEqual((k.koli, k.koli_no, k.notlar), (51, "331–381", ["last carton 10 PCS"]))

    def test_uyarilar_ve_dosyalar(self):
        self.assertTrue({("Döviz uyuşmuyor", "OZ-900"), ("Ürün kartı yok", "ZZ-000"), ("Menşe yok", "KT-050"), ("Sigorta yok", "CIF"),
                         ("Kısmi koli", "BL-310"), ("Liste altı fiyat", "HT-101")} <= self.u)
        wb = load_workbook(self.s["proforma"])
        self.assertEqual(wb.sheetnames, ["Proforma Invoice", "Packing List"])
        icerik = " ".join(str(c.value) for ws in wb for r in ws.iter_rows() for c in r if c.value is not None)
        self.assertIn("CIF Hamburg Port, Incoterms® 2020", icerik)
        self.assertNotIn("uyarı", icerik.lower())
        self.assertEqual(load_workbook(self.s["kontrol"]).sheetnames, ["Özet", "Uyarılar"])


class IncotermsTesti(unittest.TestCase):
    def calistir(self, **degisen):
        with tempfile.TemporaryDirectory() as tmp:
            s = main.calistir(ORNEK / "siparis.csv", ORNEK / "urunler.csv", bilgi(tmp, **degisen), Path(tmp))
            return s, {u["tur"] for u in s["uyarilar"]}

    def test_kurallar(self):
        s, t = self.calistir(**{"Teslim Şekli": "FOB", "Teslim Yeri": "Izmir Port"})
        self.assertIn("Gereksiz navlun", t)
        self.assertEqual(s["toplam"], s["mal"])                 # FOB'da navlun eklenmez
        s, t = self.calistir(**{"Taşıma Şekli": "Karayolu"})
        self.assertIn("Teslim şekli taşıma türüne uymuyor", t)
        s, t = self.calistir(**{"Teslim Şekli": "DAT"})
        self.assertIn("Eski teslim şekli", t)
        s, t = self.calistir(**{"Teslim Şekli": "CIP", "Taşıma Şekli": "Karayolu", "Sigorta": "310,00"})
        self.assertFalse({"Sigorta yok", "Teslim şekli taşıma türüne uymuyor", "Navlun yok"} & t)
        self.assertEqual(s["toplam"], D("82662.00") + D("2850.00") + D("310.00"))
        s, t = self.calistir(**{"Teslim Yeri": ""})
        self.assertIn("Teslim yeri yok", t)

    def test_yaziyla(self):
        self.assertEqual(main.yaziyla(D("1234.5"), "USD"), "SAY USD ONE THOUSAND TWO HUNDRED THIRTY-FOUR AND 50/100 ONLY")
        self.assertEqual(main.yaziyla(D("1000000"), "EUR"), "SAY EUR ONE MILLION AND 00/100 ONLY")
        self.assertEqual(main.yaziyla(D("0.07"), "USD"), "SAY USD ZERO AND 07/100 ONLY")
        self.assertEqual(main.yaziyla(D("110.00"), "GBP"), "SAY GBP ONE HUNDRED TEN AND 00/100 ONLY")

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", tmp]), 0)
            self.assertEqual(main.main(["--siparis", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
