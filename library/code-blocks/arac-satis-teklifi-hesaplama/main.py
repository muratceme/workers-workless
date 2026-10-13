"""
Araç Satış Teklifi Hesaplama — Workers / Workless kod bloğu
Satış › Satış Danışmanı

Liste fiyatı, kampanya, takas, peşinat, kredi ve sigorta kalemlerinden müşteriye sunulacak toplam maliyeti ve aylık taksiti hesaplar.
  Net araç fiyatı     = liste fiyatı − kampanya indirimi
  Kredi tutarı        = net araç fiyatı − takas değeri − peşinat                           (negatifse müşteriye ödenecek fark; kredi 0)
  Aylık maliyet oranı = aylık faiz % ÷ 100 × (1 + (KKDF % + BSMV %) ÷ 100)                  (--kkdf ve --bsmv; varsayılan 0)
  Aylık taksit        = kredi × oran ÷ (1 − (1 + oran)^−vade);  oran 0 ise kredi ÷ vade       (eşit taksitli)
  Ödeme planı         = her ay: faiz = kalan anapara × aylık faiz;  vergi = faiz × (KKDF + BSMV);  anapara = taksit − faiz − vergi;  son taksit kalanı kapatır
  İlk gün ödenecek    = peşinat + kasko + trafik sigortası + diğer masraflar (+ kredi yoksa kalan tutar)
  Toplam maliyet      = net araç fiyatı + kredi faizi ve vergisi + kasko + trafik sigortası + diğer masraflar
Faiz oranı, KKDF, BSMV, sigorta primleri ve masraflar sizin girdinizdir. Kod, taşıt kredisi vade ve kredilendirme oranı sınırlarını ve bankanın tahsis ücretini bilmez.
Bu çıktı bağlayıcı bir kredi teklifi değildir; kesin ödeme planını banka verir.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"no": ("teklif no", "teklif", "no"), "musteri": ("musteri", "musteri adi"), "arac": ("arac", "model", "arac modeli"), "liste": ("liste fiyati", "fiyat", "anahtar teslim fiyat", "satis fiyati"),
         "indirim": ("kampanya indirimi", "indirim", "kampanya"), "takas": ("takas degeri", "takas"), "pesinat": ("pesinat", "nakit"), "vade": ("kredi vadesi", "kredi vadesi ay", "vade", "vade ay"),
         "faiz": ("aylik faiz", "faiz", "faiz orani", "aylik faiz orani"), "kasko": ("kasko",), "trafik": ("trafik sigortasi", "trafik"), "diger": ("diger masraflar", "masraflar", "diger")}
Z = Decimal(0)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for n, r in enumerate(kayitlar(yol, SUTUN, ("liste",)), 1):
        if not metin(r["liste"]) and not metin(r["musteri"]):
            continue
        ad = metin(r["no"]) or f"Teklif {n}"
        d = {a: para(r[a]) if metin(r[a]) else Z for a in ("liste", "indirim", "takas", "pesinat", "faiz", "kasko", "trafik", "diger")}
        vade = para(r["vade"]) if metin(r["vade"]) else Z
        bozuk = [a for a, v in d.items() if v is None or v < 0]
        if bozuk or not d["liste"] or vade is None or vade < 0 or vade != vade.to_integral_value() or d["indirim"] > d["liste"]:
            uy.append(uyari("Yüksek", "Teklif okunamadı", f"{ad} (satır {r['_satir']})", f"Liste fiyatı pozitif, diğer tutarlar sıfır veya pozitif, vade tam ay olmalı; indirim liste fiyatını aşamaz"
                            + (f" (sorunlu alan: {', '.join(bozuk)})" if bozuk else "") + ". Teklif hesaplanmadı."))
            continue
        l.append({"no": ad, "musteri": metin(r["musteri"]), "arac": metin(r["arac"]), "vade": int(vade)} | d)
    return l, uy


def taksit_hesapla(kredi: Decimal, faiz: Decimal, vade: int, vergi: Decimal = Z) -> tuple[Decimal, list[dict]]:
    """Eşit taksit ve ödeme planı. faiz: aylık %, vergi: KKDF + BSMV toplamı %."""
    i, v = faiz / 100, vergi / 100
    oran = i * (1 + v)
    taksit = yuvarla(kredi * oran / (1 - (1 + oran) ** -vade)) if oran else yuvarla(kredi / vade)
    plan, kalan = [], kredi
    for ay in range(1, vade + 1):
        f, vg = yuvarla(kalan * i), yuvarla(kalan * i * v)
        ana = kalan if ay == vade else min(taksit - f - vg, kalan)
        kalan -= ana
        plan.append({"ay": ay, "taksit": ana + f + vg, "faiz": f, "vergi": vg, "anapara": ana, "kalan": kalan})
    return taksit, plan


def hesapla(teklifler: list[dict], kkdf: Decimal = Z, bsmv: Decimal = Z) -> dict:
    if not teklifler:
        raise ValueError("Hesaplanacak teklif yok.")
    uy, liste, planlar = [], [], []
    for t in teklifler:
        net = t["liste"] - t["indirim"]
        fark = net - t["takas"] - t["pesinat"]
        kredi, iade, ek = max(fark, Z), max(-fark, Z), t["kasko"] + t["trafik"] + t["diger"]
        taksit = geri = faiz = vergi = None
        nakit_kalan = Z
        if kredi and t["vade"]:
            taksit, plan = taksit_hesapla(kredi, t["faiz"], t["vade"], kkdf + bsmv)
            geri, faiz, vergi = sum(p["taksit"] for p in plan), sum(p["faiz"] for p in plan), sum(p["vergi"] for p in plan)
            planlar += [{"no": t["no"], "musteri": t["musteri"]} | p for p in plan]
            if not t["faiz"]:
                uy.append(uyari("Bilgi", "Faizsiz kredi", t["no"], "Aylık faiz 0 girildi; taksit anaparanın vadeye bölünmesiyle hesaplandı. Kampanyalı sıfır faiz değilse oranı kontrol edin."))
        elif kredi:
            nakit_kalan = kredi
            uy.append(uyari("Orta", "Kalan tutar için vade yok", t["no"], f"Takas ve peşinattan sonra {kredi:,.2f} kalıyor ama kredi vadesi girilmemiş. Tutar teslimde nakit ödenecek sayıldı."))
        if iade:
            uy.append(uyari("Orta", "Takas ve peşinat araç fiyatını aşıyor", t["no"], f"Takas {t['takas']:,.2f} + peşinat {t['pesinat']:,.2f}, net araç fiyatından {iade:,.2f} fazla. Peşinatı düşürün veya farkın müşteriye ödenmesini planlayın."))
        if t["faiz"] > 10:
            uy.append(uyari("Orta", "Aylık faiz çok yüksek görünüyor", t["no"], f"Aylık faiz %{t['faiz'].normalize():f} girildi. Yıllık oran yazılmış olabilir; sütun aylık oran bekler."))
        maliyet = net + (faiz or Z) + (vergi or Z) + ek
        liste.append({"no": t["no"], "musteri": t["musteri"], "arac": t["arac"], "liste": t["liste"], "indirim": t["indirim"], "net": net, "takas": t["takas"], "pesinat": t["pesinat"], "kredi": kredi if t["vade"] else Z, "iade": iade,
                      "kredi_oran": yuvarla(kredi / net * 100) if t["vade"] and net else Z, "vade": t["vade"] if kredi else 0, "faiz_y": t["faiz"], "taksit": taksit, "geri": geri, "faiz": faiz, "vergi": vergi, "ek": ek,
                      "ilk_gun": t["pesinat"] + ek + nakit_kalan - min(iade, t["pesinat"]), "maliyet": maliyet})
    return {"teklifler": liste, "planlar": planlar, "ayar": (kkdf, bsmv), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    kkdf, bsmv = r["ayar"]
    ws = wb.active
    ws.title = "Teklifler"
    _baslik(ws, ["Teklif No", "Müşteri", "Araç", "Liste Fiyatı", "Kampanya İndirimi", "Net Araç Fiyatı", "Takas", "Peşinat", "Kredi Tutarı", "Kredi Oranı %", "Vade (ay)", "Aylık Faiz %", "Aylık Taksit", "Toplam Geri Ödeme", "Faiz", "KKDF + BSMV",
                 "Sigorta ve Masraflar", "İlk Gün Ödenecek", "Toplam Maliyet"], (12, 26, 26, 15, 13, 15, 14, 14, 15, 9, 8, 9, 14, 16, 14, 13, 14, 15, 16), "C2")
    for t in r["teklifler"]:
        _satir(ws, [t["no"], t["musteri"], t["arac"], t["liste"], t["indirim"] or None, t["net"], t["takas"] or None, t["pesinat"] or None, t["kredi"] or None, t["kredi_oran"] or None, t["vade"] or None, t["faiz_y"] if t["taksit"] else None,
                    t["taksit"], t["geri"], t["faiz"], t["vergi"], t["ek"] or None, t["ilk_gun"], t["maliyet"]], para_s=(4, 5, 6, 7, 8, 9, 10, 12, 13, 14, 15, 16, 17, 18, 19), dolgu={13: SARI if t["taksit"] else None, 19: MAVI})
    ws.append([])
    ws.append([f"KKDF %{kkdf.normalize():f}, BSMV %{bsmv.normalize():f} ile hesaplandı. Bağlayıcı kredi teklifi değildir; kesin ödeme planını, tahsis ücretini ve kredilendirme sınırlarını banka belirler."])
    ws = wb.create_sheet("Ödeme Planı")
    _baslik(ws, ["Teklif No", "Müşteri", "Ay", "Taksit", "Faiz", "KKDF + BSMV", "Anapara", "Kalan Anapara"], (12, 26, 6, 14, 13, 13, 14, 15))
    for p in r["planlar"]:
        _satir(ws, [p["no"], p["musteri"], p["ay"], p["taksit"], p["faiz"], p["vergi"] or None, p["anapara"], p["kalan"]], para_s=(4, 5, 6, 7, 8))
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, kkdf: Decimal = Z, bsmv: Decimal = Z) -> dict:
    teklifler, uy = oku(yol)
    r = hesapla(teklifler, kkdf, bsmv)
    uy += r["uyarilar"]
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Araç satış teklifinde kredi taksitini, ilk gün ödenecek tutarı ve toplam maliyeti hesaplar.")
    p.add_argument("--teklifler", type=Path, default=ORNEK / "teklifler.csv", help="Teklif No, Müşteri, Araç, Liste Fiyatı, Kampanya İndirimi, Takas Değeri, Peşinat, Kredi Vadesi (ay), Aylık Faiz %%, Kasko, Trafik Sigortası, Diğer Masraflar")
    p.add_argument("--kkdf", default="0", help="Kredi faizi üzerinden KKDF oranı, %% (varsayılan 0)")
    p.add_argument("--bsmv", default="0", help="Kredi faizi üzerinden BSMV oranı, %% (varsayılan 0)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "arac_satis_teklifi.xlsx")
    a = p.parse_args(argv)
    k, b = para(a.kkdf), para(a.bsmv)
    if k is None or b is None or not 0 <= k < 100 or not 0 <= b < 100:
        print("[X] --kkdf ve --bsmv 0 ile 100 arasında olmalı")
        return 2
    if not a.teklifler.exists():
        print(f"[X] Dosya bulunamadı: {a.teklifler}")
        return 1
    try:
        r = calistir(a.teklifler, a.cikti, k, b)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    for t in r["teklifler"]:
        kredi = f" · kredi {t['kredi']:,.2f}, {t['vade']} ay × {t['taksit']:,.2f}" if t["taksit"] else ""
        print(f"[OK] {t['no']} {t['musteri']}: net {t['net']:,.2f}{kredi} · ilk gün {t['ilk_gun']:,.2f} · toplam maliyet {t['maliyet']:,.2f}")
    for u in r["uyarilar"]:
        if u["onem"] != "Bilgi":
            print(f"[!] {u['tur']}: {u['kim']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
