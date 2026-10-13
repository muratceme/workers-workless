import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def t(liste, indirim=0, takas=0, pesinat=0, vade=0, faiz=0, kasko=0, trafik=0, diger=0, no="T1"):
    return {"no": no, "musteri": "M", "arac": "A", "liste": D(liste), "indirim": D(indirim), "takas": D(takas), "pesinat": D(pesinat), "vade": vade, "faiz": D(faiz), "kasko": D(kasko), "trafik": D(trafik), "diger": D(diger)}


class TaksitTesti(unittest.TestCase):
    def test_esit_taksit_ve_plan(self):
        taksit, plan = main.taksit_hesapla(D(10000), D(10), 2)
        self.assertEqual((taksit, [(p["ay"], p["taksit"], p["faiz"], p["vergi"], p["anapara"], p["kalan"]) for p in plan]),
                         (D("5761.90"), [(1, D("5761.90"), D("1000.00"), D("0.00"), D("4761.90"), D("5238.10")), (2, D("5761.91"), D("523.81"), D("0.00"), D("5238.10"), D("0.00"))]))
        taksit, plan = main.taksit_hesapla(D(10000), D(10), 2, D(30))
        self.assertEqual((taksit, [(p["taksit"], p["faiz"], p["vergi"], p["anapara"], p["kalan"]) for p in plan]),
                         (D("5994.84"), [(D("5994.84"), D("1000.00"), D("300.00"), D("4694.84"), D("5305.16")), (D("5994.83"), D("530.52"), D("159.15"), D("5305.16"), D("0.00"))]))
        taksit, plan = main.taksit_hesapla(D(120000), D(0), 12)
        self.assertEqual((taksit, {p["taksit"] for p in plan}, plan[-1]["kalan"]), (D("10000.00"), {D("10000.00")}, D("0.00")))
        taksit, plan = main.taksit_hesapla(D(400000), D(3), 12)
        self.assertEqual((taksit, plan[-1]["kalan"], sum(p["anapara"] for p in plan), abs(plan[-1]["taksit"] - taksit) < D("0.10")), (D("40184.83"), D("0.00"), D("400000.00"), True))


class TeklifTesti(unittest.TestCase):
    def test_kredili(self):
        r = main.hesapla([t(1000000, 50000, 300000, 250000, 12, 3, 30000, 10000, 5000)])
        x = r["teklifler"][0]
        self.assertEqual((x["net"], x["kredi"], x["kredi_oran"], x["vade"], x["taksit"], x["ek"], x["ilk_gun"], x["iade"], x["vergi"], r["uyarilar"], len(r["planlar"])),
                         (D(950000), D(400000), D("42.11"), 12, D("40184.83"), D(45000), D(295000), D(0), D("0.00"), [], 12))
        self.assertEqual((x["geri"], x["maliyet"], abs(x["geri"] - D("482217.96")) < D("0.10")), (D(400000) + x["faiz"], D(995000) + x["faiz"], True))
        v = main.hesapla([t(1000000, 50000, 300000, 250000, 12, 3)], D(15), D(15))["teklifler"][0]
        self.assertEqual((v["taksit"] > x["taksit"], v["vergi"] > 0, v["geri"], v["maliyet"]), (True, True, D(400000) + v["faiz"] + v["vergi"], D(950000) + v["faiz"] + v["vergi"]))

    def test_nakit_iade_ve_uyarilar(self):
        r = main.hesapla([t(500000, 0, 100000, 100000, kasko=20000, no="N"), t(800000, pesinat=800000, no="P"), t(500000, 0, 400000, 150000, 12, 3, no="I"), t(100000, vade=10, no="F"), t(100000, vade=12, faiz=45, no="Y")])
        d = {x["no"]: x for x in r["teklifler"]}
        self.assertEqual([(d[k]["kredi"], d[k]["taksit"], d[k]["ilk_gun"], d[k]["maliyet"], d[k]["iade"]) for k in ("N", "P", "I", "F")],
                         [(D(0), None, D(420000), D(520000), D(0)), (D(0), None, D(800000), D(800000), D(0)), (D(0), None, D(100000), D(500000), D(50000)), (D(100000), D("10000.00"), D(0), D("100000.00"), D(0))])
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Aylık faiz çok yüksek görünüyor", "Y"), ("Faizsiz kredi", "F"), ("Kalan tutar için vade yok", "N"), ("Takas ve peşinat araç fiyatını aşıyor", "I")])
        with self.assertRaises(ValueError):
            main.hesapla([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "a.xlsx"
            r = main.calistir(ORNEK / "teklifler.csv", cikti)
            self.assertEqual([(x["no"], x["net"], x["kredi"], x["taksit"], x["ilk_gun"]) for x in r["teklifler"]],
                             [("T-001", D("950000.00"), D("400000.00"), D("40184.83"), D("295000.00")), ("T-002", D("800000.00"), D(0), None, D("838000.00")), ("T-003", D("1500000.00"), D(0), None, D("962000.00"))])
            self.assertEqual(([(u["tur"], u["kim"]) for u in r["uyarilar"]], len(r["planlar"])), ([("Kalan tutar için vade yok", "T-003")], 12))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Teklifler", "Ödeme Planı", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "t.csv"
            y.write_text("Müşteri;Fiyat;İndirim;Vade;Faiz\nA;1.250.000,50;;24;2,5\nB;0;;;\nC;100;200;;\nD;100;;2,5;\nE;100;-5;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["musteri"], x["liste"], x["vade"], x["faiz"], x["kasko"]) for x in l], [u["tur"] for u in uy]), ([("Teklif 1", "A", D("1250000.50"), 24, D("2.5"), D(0))], ["Teklif okunamadı"] * 4))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--teklifler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--kkdf", "15", "--bsmv", "15"]), 0)
            for ek in (["--kkdf", "x"], ["--bsmv", "100"]):
                self.assertEqual(main.main(arg + ek), 2)
            self.assertEqual(main.main(["--teklifler", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Müşteri;Vade\nA;1\n", "Müşteri;Fiyat\nA;x\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
