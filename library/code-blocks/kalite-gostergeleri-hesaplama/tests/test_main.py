import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def tn(gosterge="G", pay="Olay", payda="Gün", carpan=1000, hedef=1, yon="Düşük iyi"):
    return {"gosterge": gosterge, "pay": pay, "payda": payda, "carpan": D(carpan), "hedef": None if hedef is None else D(hedef), "yon": yon, "birim": ""}


def vr(**donemler):
    """vr(d09={"A": (3, 2000)}) → 2026-09 döneminde A biriminde 3 olay, 2000 gün"""
    veri, adlar = {}, {}
    for d, birimler in donemler.items():
        for b, (olay, gun) in birimler.items():
            adlar[b.lower()] = b
            veri[f"2026-{d[1:]}", b.lower()] = {k: D(v) for k, v in (("olay", olay), ("gun", gun)) if v is not None}
    return veri, adlar


class HesapTesti(unittest.TestCase):
    def test_deger_ve_hedef(self):
        r = main.hesapla([tn()], *vr(d09={"A": (3, 2000), "B": (0, 500)}))
        o = r["ozet"][0]
        self.assertEqual((o["pay"], o["payda"], o["deger"], o["durum"], o["onceki"], o["egilim"]), (D(3), D(2500), D("1.20"), "Hedef dışı", None, ""))      # toplamdan: 3 ÷ 2500, oranların ortalaması değil
        self.assertEqual([(b["birim"], b["deger"], b["durum"]) for b in r["birimler"]], [("A", D("1.50"), "Hedef dışı"), ("B", D("0.00"), "Hedefte")])
        self.assertEqual([main.hesapla([tn(hedef=h)], *vr(d09={"A": (2, 2000)}))["ozet"][0]["durum"] for h in (1, "0.99", None)], ["Hedefte", "Hedef dışı", ""])      # tam hedefte: Hedefte
        self.assertEqual([main.hesapla([tn(carpan=100, hedef=80, yon="Yüksek iyi")], *vr(d09={"A": (p, 100)}))["ozet"][0]["durum"] for p in (80, 79)], ["Hedefte", "Hedef dışı"])
        o = main.hesapla([tn()], *vr(d09={"A": (1, 0)}))["ozet"][0]
        self.assertEqual((o["deger"], o["durum"]), (None, ""))

    def test_egilim_ve_donem_secimi(self):
        v = vr(d07={"A": (4, 1000)}, d08={"A": (2, 1000)}, d09={"A": (3, 1000)})
        r = main.hesapla([tn(), tn("Y", yon="Yüksek iyi")], *v)
        self.assertEqual((r["rapor"], r["onceki"], [(o["deger"], o["onceki"], o["fark"], o["egilim"]) for o in r["ozet"]]),
                         ("2026-09", "2026-08", [(D("3.00"), D("2.00"), D("1.00"), "Kötüleşti"), (D("3.00"), D("2.00"), D("1.00"), "İyileşti")]))
        r = main.hesapla([tn()], *v, donem="2026-08")
        self.assertEqual((r["onceki"], r["ozet"][0]["egilim"], r["egilim"][0]["seri"]), ("2026-07", "İyileşti", {"2026-07": D("4.00"), "2026-08": D("2.00"), "2026-09": D("3.00")}))
        self.assertEqual(main.hesapla([tn()], *vr(d08={"A": (2, 1000)}, d09={"A": (2, 1000)}))["ozet"][0]["egilim"], "Değişmedi")
        self.assertEqual(main.hesapla([tn()], *v, donem="2026-07")["onceki"], None)
        for kotu in (({}, {}, None), (*v, "2026-01")):
            with self.assertRaises(ValueError):
                main.hesapla([tn()], *kotu)

    def test_eksik_veri(self):
        r = main.hesapla([tn()], *vr(d09={"A": (3, 2000), "B": (None, 500), "C": (None, None)}))
        self.assertEqual((r["ozet"][0]["deger"], r["ozet"][0]["eksik"]), (D("1.50"), 2))                              # yalnız tam verisi olan A toplamda
        self.assertEqual([(b["birim"], b["deger"]) for b in r["birimler"]], [("A", D("1.50")), ("B", None)])           # hiç verisi olmayan C listelenmez
        self.assertEqual(r["eksik"], [{"donem": "2026-09", "birim": "B", "gosterge": "G", "veri": "Olay"}])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "k.xlsx"
            r = main.calistir(ORNEK / "gosterge_verileri.csv", ORNEK / "gosterge_tanimlari.csv", cikti)
            self.assertEqual([(o["gosterge"].split(" (")[0], o["pay"], o["payda"], o["deger"], o["durum"], o["onceki"], o["egilim"]) for o in r["ozet"]],
                             [("Düşme oranı", D(3), D(2500), D("1.20"), "Hedef dışı", D("0.84"), "Kötüleşti"), ("Bası yarası oranı", D(7), D(460), D("1.52"), "Hedefte", D("1.84"), "İyileşti"),
                              ("Enfeksiyon hızı", D(7), D(2500), D("2.80"), "Hedefte", D("2.63"), "Kötüleşti")])
            self.assertEqual([(b["birim"], b["deger"], b["durum"]) for b in r["birimler"]],
                             [("Dahiliye Servisi", D("1.50"), "Hedef dışı"), ("Yoğun Bakım", D("0.00"), "Hedefte"), ("Dahiliye Servisi", D("1.00"), "Hedefte"), ("Yoğun Bakım", D("5.00"), "Hedef dışı"),
                              ("Dahiliye Servisi", D("2.00"), "Hedefte"), ("Yoğun Bakım", D("6.00"), "Hedef dışı")])
            self.assertEqual((r["say"], [(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]]),
                             ({"gosterge": 3, "hedefte": 2, "hedef_disi": 1, "hesaplanamayan": 0, "birim_hedef_disi": 3}, [("Bilgi", "Gösterge verisi eksik", "2026-08 · Yoğun Bakım")]))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Gösterge Özeti", "Birim Bazında", "Dönem Eğilimi", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            v, t = Path(tmp) / "v.csv", Path(tmp) / "t.csv"
            t.write_text("Gösterge;Pay;Payda;Çarpan;Hedef;Yön\nA;Olay;Gün;;düşük;\nB;Uygun;Gözlem;100;80;YÜKSEK İYİ\nC;;Gün;100;;\nD;Olay;Gün;0;;\nE;Yok;Gün;;;\n", encoding="utf-8")
            v.write_text("Ay;Servis;Veri Adı;Sayı\nEylül 2026;S1;OLAY;2\n09.2026;S1;olay;1\n2026-09;S1;Gün;1.000\n2026-09;;Gün;500\n3. çeyrek;S1;Gün;5\n2026-09;S1;Uygun;çok\n", encoding="utf-8")
            tl, uy = main.tanim_oku(t)
            self.assertEqual([(x["gosterge"], x["carpan"], x["hedef"], x["yon"]) for x in tl], [("A", D(100), None, "Düşük iyi"), ("B", D(100), D(80), "Yüksek iyi"), ("E", D(100), None, "Düşük iyi")])
            self.assertEqual(sorted(u["tur"] for u in uy), ["Gösterge tanımı eksik", "Gösterge tanımı eksik", "Hedef okunamadı"])
            veri, adlar, uy = main.oku(v)
            self.assertEqual((veri, adlar, len(uy)), ({("2026-09", "s1"): {"olay": D(3), "gun": D(1000)}, ("2026-09", "kurum"): {"gun": D(500)}}, {"s1": "S1", "kurum": "Kurum"}, 2))
            c = str(Path(tmp) / "x.xlsx")
            r = main.calistir(v, t, Path(c))
            self.assertEqual([(o["gosterge"], o["deger"]) for o in r["ozet"]], [("A", D("0.30")), ("B", None), ("E", None)])
            self.assertEqual(sorted({u["tur"] for u in r["uyarilar"]}), ["Dönem veya değer okunamadı", "Gösterge tanımı eksik", "Gösterge verisi eksik", "Hedef okunamadı", "Tanımdaki veri hiç girilmemiş"])
            arg = ["--veriler", str(v), "--tanimlar", str(t), "--cikti", c]
            self.assertEqual(main.main(arg + ["--donem", "Eylül 2026"]), 0)
            self.assertEqual(main.main(arg + ["--donem", "2025-01"]), 1)
            self.assertEqual(main.main(arg + ["--donem", "geçen ay"]), 2)
            self.assertEqual(main.main(arg[:3] + [str(Path(tmp) / "yok.csv")]), 1)
            t.write_text("Gösterge;Pay;Payda\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
