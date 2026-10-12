"""
Kalite Göstergeleri Hesaplama — Workers / Workless kod bloğu
Kalite Yönetimi › Kalite Birim Sorumlusu

Kayıtlardan SKS kalite göstergelerini (düşme, bası yarası, enfeksiyon oranı vb.) formüllerine göre dönemsel hesaplar.
  Gösterge değeri = pay verisi ÷ payda verisi × çarpan           (gösterge tanım dosyasından: hangi veri paya, hangisi paydaya girer, çarpan 100 / 1000 ...)
  Kurum toplamı   = birimlerin pay toplamı ÷ payda toplamı × çarpan   (birim oranlarının ortalaması değildir; pay veya paydası eksik birim toplamın dışında kalır ve uyarılır)
  Hedef durumu    = yön "Düşük iyi" ise değer ≤ hedef: Hedefte;  "Yüksek iyi" ise değer ≥ hedef: Hedefte;  aksi halde Hedef dışı
  Eğilim          = rapor döneminin kurum değeri önceki dönemle karşılaştırılır; yöne göre İyileşti / Kötüleşti / Değişmedi
Formüller koda gömülü değildir: gösterge tanımlarını kendi gösterge kartlarınıza (SKS gösterge yönetimi rehberi, kurum prosedürü) göre siz yazarsınız.
Örnek tanım dosyasındaki formül ve hedefler örnektir, doğrulanmamıştır. Hasta bilgisi gerekmez; yalnız sayılar girilir.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


TANIM = {"gosterge": ("gosterge", "gosterge adi"), "pay": ("pay", "pay verisi"), "payda": ("payda", "payda verisi"), "carpan": ("carpan", "katsayi"), "hedef": ("hedef", "hedef deger"),
         "yon": ("yon", "iyi yon", "hedef yonu"), "birim": ("olcu birimi", "birim", "olcu")}
VERI = {"donem": ("donem", "ay"), "birim": ("birim", "servis", "bolum", "klinik"), "veri": ("veri", "veri adi", "kalem", "gosterge verisi"), "deger": ("deger", "sayi", "adet", "miktar")}
AYLAR = ("ocak", "subat", "mart", "nisan", "mayis", "haziran", "temmuz", "agustos", "eylul", "ekim", "kasim", "aralik")
Z = Decimal(0)


def donem_coz(x) -> tuple[int, int] | None:
    """2026-08 · 08.2026 · 2026/8 · Ağustos 2026 · tarih → (yıl, ay)"""
    if isinstance(x, (date, datetime)):
        return x.year, x.month
    k = katla(x)
    m = re.fullmatch(r"(\d{4}) (\d{1,2})", k) or re.fullmatch(r"(\d{4}) (\d{1,2}) \d{1,2}.*", k)
    if m and 1 <= int(m[2]) <= 12:
        return int(m[1]), int(m[2])
    m = re.fullmatch(r"(?:\d{1,2} )?(\d{1,2}) (\d{4})", k)
    if m and 1 <= int(m[1]) <= 12:
        return int(m[2]), int(m[1])
    m = re.fullmatch(r"([a-z]+) (\d{4})", k)
    if m and m[1] in AYLAR:
        return int(m[2]), AYLAR.index(m[1]) + 1
    return None


def donem_adi(x) -> str | None:
    d = donem_coz(x)
    return f"{d[0]}-{d[1]:02d}" if d else None


def tanim_oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, TANIM, ("gosterge", "pay", "payda")):
        g, pay, payda = metin(r["gosterge"]), metin(r["pay"]), metin(r["payda"])
        if not g:
            continue
        carpan, hedef = para(r["carpan"]) if metin(r["carpan"]) else Decimal(100), para(r["hedef"])
        if not pay or not payda or carpan is None or carpan <= 0:
            uy.append(uyari("Yüksek", "Gösterge tanımı eksik", g, f"{yol.name} satır {r['_satir']}: pay, payda veya çarpan okunamadı. Gösterge hesaplanmadı."))
            continue
        if hedef is None and metin(r["hedef"]):
            uy.append(uyari("Orta", "Hedef okunamadı", g, f"'{metin(r['hedef'])}'. Hedef karşılaştırması yapılmadı."))
        l.append({"gosterge": g, "pay": pay, "payda": payda, "carpan": carpan, "hedef": hedef, "yon": "Yüksek iyi" if "yuksek" in katla(r["yon"]) else "Düşük iyi", "birim": metin(r["birim"])})
    if not l:
        raise ValueError(f"{yol.name}: hesaplanabilir gösterge tanımı yok.")
    return l, uy


def oku(yol: Path) -> tuple[dict, dict, list[dict]]:
    """veri = {(dönem, birim anahtarı): {veri anahtarı: toplam değer}}, adlar = {birim anahtarı: ad}"""
    uy, veri, adlar = [], defaultdict(lambda: defaultdict(lambda: Z)), {}
    for r in kayitlar(yol, VERI, ("donem", "veri", "deger")):
        v = metin(r["veri"])
        if not v:
            continue
        kim = f"{v} (satır {r['_satir']})"
        d, deger = donem_adi(r["donem"]), para(r["deger"])
        if d is None or deger is None or deger < 0:
            uy.append(uyari("Yüksek", "Dönem veya değer okunamadı", kim, f"Dönem '{metin(r['donem'])}', değer '{metin(r['deger'])}'. Satır alınmadı. Dönem biçimi: 2026-09, 09.2026 veya Eylül 2026."))
            continue
        b = metin(r["birim"]) or "Kurum"
        adlar.setdefault(katla(b), b)
        veri[d, katla(b)][katla(v)] += deger
    return {k: dict(v) for k, v in veri.items()}, adlar, uy


def hesapla(tanimlar: list[dict], veri: dict, adlar: dict, donem: str | None = None) -> dict:
    if not veri:
        raise ValueError("Hesaplanacak veri yok.")
    donemler = sorted({d for d, _ in veri})
    rapor = donem or donemler[-1]
    if rapor not in donemler:
        raise ValueError(f"{rapor} dönemine ait veri yok. Verideki dönemler: {', '.join(donemler)}.")
    onceki = donemler[donemler.index(rapor) - 1] if donemler.index(rapor) else None

    def oran(t: dict, kaynaklar: list[dict]) -> dict:
        p, q = katla(t["pay"]), katla(t["payda"])
        tam = [k for k in kaynaklar if p in k and q in k]
        pay, payda = sum((k[p] for k in tam), Z), sum((k[q] for k in tam), Z)
        v = yuvarla(pay / payda * t["carpan"]) if tam and payda else None
        durum = ""
        if v is not None and t["hedef"] is not None:
            durum = "Hedefte" if (v <= t["hedef"] if t["yon"] == "Düşük iyi" else v >= t["hedef"]) else "Hedef dışı"
        return {"pay": pay if tam else None, "payda": payda if tam else None, "deger": v, "durum": durum, "eksik": len(kaynaklar) - len(tam), "kaynak": len(kaynaklar)}

    ozet, birimler, egilim, eksik = [], [], [], []
    for t in tanimlar:
        seri = {d: oran(t, [v for (dd, _), v in veri.items() if dd == d]) for d in donemler}
        bu, once = seri[rapor], seri[onceki] if onceki else None
        fark = bu["deger"] - once["deger"] if once and bu["deger"] is not None and once["deger"] is not None else None
        yon = "" if fark is None else "Değişmedi" if fark == 0 else "İyileşti" if (fark < 0) == (t["yon"] == "Düşük iyi") else "Kötüleşti"
        ozet.append({"gosterge": t["gosterge"], "pay_ad": t["pay"], "payda_ad": t["payda"], "carpan": t["carpan"], "hedef": t["hedef"], "yon": t["yon"], "birim": t["birim"]} | bu
                    | {"onceki": once["deger"] if once else None, "fark": fark, "egilim": yon})
        egilim.append({"gosterge": t["gosterge"], "seri": {d: seri[d]["deger"] for d in donemler}, "hedef": t["hedef"], "yon": t["yon"]})
        for (d, b), v in sorted(veri.items(), key=lambda x: (x[0][0], adlar[x[0][1]])):
            var = [a for a in (t["pay"], t["payda"]) if katla(a) in v]
            if not var:                                # birim bu göstergeyi hiç izlemiyor
                continue
            if d == rapor:
                birimler.append({"gosterge": t["gosterge"], "birim": adlar[b], "hedef": t["hedef"], "olcu": t["birim"]} | oran(t, [v]))
            if len(var) == 1:
                eksik.append({"donem": d, "birim": adlar[b], "gosterge": t["gosterge"], "veri": t["payda"] if var[0] == t["pay"] else t["pay"]})
    return {"rapor": rapor, "onceki": onceki, "donemler": donemler, "ozet": ozet, "birimler": birimler, "egilim": egilim, "eksik": eksik,
            "say": {"gosterge": len(ozet), "hedefte": sum(o["durum"] == "Hedefte" for o in ozet), "hedef_disi": sum(o["durum"] == "Hedef dışı" for o in ozet),
                    "hesaplanamayan": sum(o["deger"] is None for o in ozet), "birim_hedef_disi": sum(b["durum"] == "Hedef dışı" for b in birimler)}}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    renk = {"Hedefte": YESIL, "Hedef dışı": KIRMIZI}
    ws = wb.active
    ws.title = "Gösterge Özeti"
    _baslik(ws, ["Gösterge", "Pay Verisi", "Payda Verisi", "Çarpan", "Pay", "Payda", "Değer", "Ölçü Birimi", "Hedef", "Yön", "Durum", f"Önceki Dönem ({r['onceki'] or '—'})", "Fark", "Eğilim",
                 "İyileştirme Faaliyeti (siz doldurun)"], (34, 32, 28, 8, 10, 10, 10, 14, 9, 11, 11, 12, 9, 11, 36))
    for o in r["ozet"]:
        _satir(ws, [o["gosterge"], o["pay_ad"], o["payda_ad"], o["carpan"], o["pay"], o["payda"], o["deger"] if o["deger"] is not None else "hesaplanamadı", o["birim"], o["hedef"], o["yon"], o["durum"],
                    o["onceki"], o["fark"], o["egilim"], ""], para_s=(7, 9, 12, 13), dolgu={11: renk.get(o["durum"]), 14: {"İyileşti": YESIL, "Kötüleşti": SARI}.get(o["egilim"])}, sar=(1, 2, 3))
    ws.append([])
    ws.append([f"Rapor dönemi {r['rapor']}. Değer = pay ÷ payda × çarpan; kurum değeri birimlerin pay ve payda toplamlarından hesaplanır. Formül ve hedefler gösterge tanım dosyanızdan gelir; "
               "kendi gösterge kartlarınızla karşılaştırın."])
    ws = wb.create_sheet("Birim Bazında")
    _baslik(ws, ["Gösterge", "Birim", "Pay", "Payda", "Değer", "Ölçü Birimi", "Hedef", "Durum"], (34, 26, 10, 10, 10, 14, 9, 14))
    for b in r["birimler"]:
        _satir(ws, [b["gosterge"], b["birim"], b["pay"], b["payda"], b["deger"] if b["deger"] is not None else "veri eksik", b["olcu"], b["hedef"], b["durum"]], para_s=(5, 7), dolgu={8: renk.get(b["durum"])})
    ws = wb.create_sheet("Dönem Eğilimi")
    _baslik(ws, ["Gösterge", "Hedef", "Yön", *r["donemler"]], (34, 9, 11, *[11] * len(r["donemler"])))
    for e in r["egilim"]:
        _satir(ws, [e["gosterge"], e["hedef"], e["yon"], *[e["seri"][d] for d in r["donemler"]]], para_s=(2, *range(4, 4 + len(r["donemler"]))))
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(veri_yol: Path, tanim_yol: Path, cikti: Path, donem: str | None = None) -> dict:
    tanimlar, uy = tanim_oku(tanim_yol)
    veri, adlar, uy2 = oku(veri_yol)
    r = hesapla(tanimlar, veri, adlar, donem)
    uy += uy2
    girilen = {a for v in veri.values() for a in v}
    for t in tanimlar:
        for a in (t["pay"], t["payda"]):
            if katla(a) not in girilen:
                uy.append(uyari("Yüksek", "Tanımdaki veri hiç girilmemiş", t["gosterge"], f"'{a}' veri dosyasında hiç geçmiyor. Veri adının iki dosyada aynı yazıldığını kontrol edin."))
    for e in r["eksik"]:
        uy.append(uyari("Orta" if e["donem"] == r["rapor"] else "Bilgi", "Gösterge verisi eksik", f"{e['donem']} · {e['birim']}", f"{e['gosterge']}: '{e['veri']}' girilmemiş. Birim bu göstergede kurum toplamına katılmadı."))
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Gösterge tanımlarına göre kalite göstergelerini dönemsel hesaplar.")
    p.add_argument("--veriler", type=Path, default=ORNEK / "gosterge_verileri.csv", help="Dönem, Birim, Veri, Değer (her satır bir sayım)")
    p.add_argument("--tanimlar", type=Path, default=ORNEK / "gosterge_tanimlari.csv", help="Gösterge, Pay, Payda, Çarpan, Hedef, Yön (Düşük iyi / Yüksek iyi), Ölçü Birimi")
    p.add_argument("--donem", help="Rapor dönemi (ör. 2026-09); verilmezse verideki son dönem")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kalite_gostergeleri.xlsx")
    a = p.parse_args(argv)
    donem = donem_adi(a.donem) if a.donem else None
    if a.donem and donem is None:
        print("[X] --donem 2026-09, 09.2026 veya 'Eylül 2026' biçiminde olmalı")
        return 2
    for y in (a.veriler, a.tanimlar):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.veriler, a.tanimlar, a.cikti, donem)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    s = r["say"]
    print(f"[OK] Dönem {r['rapor']} · {s['gosterge']} gösterge · {s['hedefte']} hedefte, {s['hedef_disi']} hedef dışı, {s['hesaplanamayan']} hesaplanamadı · birim bazında {s['birim_hedef_disi']} hedef dışı")
    for o in r["ozet"]:
        print(f"[{'!' if o['durum'] == 'Hedef dışı' else 'OK'}] {o['gosterge']}: {o['deger'] if o['deger'] is not None else 'hesaplanamadı'} {o['birim']}" + (f" (hedef {o['hedef']:f}, {o['durum']})" if o["durum"] else "")
              + (f" · {o['egilim']}" if o["egilim"] else ""))
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
