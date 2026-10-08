"""
Ücret Bandı Karşılaştırması — Workers / Workless kod bloğu
İnsan Kaynakları › Ücretlendirme ve Yan Haklar Uzmanı

İç ücret verisini kademe / pozisyon bazında ücret araştırması bantlarıyla kıyaslar:
  - Bant: önce pozisyona özel bant, yoksa kademe bandı (Alt / Orta / Üst; ör. P25 / P50 / P75).
  - Kısmi süreli çalışanlarda ücret tam zamanlı karşılığına çevrilir (ücret ÷ çalışma oranı).
  - Karşılaştırma oranı (compa-ratio) = ücret ÷ orta nokta; bant içi konum = (ücret − alt) ÷ (üst − alt).
  - Bant altı (Yüksek), bant üstü (Orta), yasal asgari ücretin altı (Yüksek; tr_parametreler.json, yıl bazında).
  - Ücret sıkışması: aynı pozisyonda en az 2 yıl daha kıdemli ve performansı eşit / daha iyi çalışan, yeni gelenin
    ücretinin %95'inden az alıyor.
  - Bütçe etkisi: bant altındakileri alt sınıra çekmenin aylık ve yıllık brüt maliyeti.
  - Cinsiyete göre kademe ortalama karşılaştırma oranı (grupta en az 2 kişi varsa; %5 üstü fark incelenmeli).
Rapor: çalışan listesi, bant dışı / uyarılar, kademe ve departman özeti, cinsiyet karşılaştırması. İnternete
bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: 28 çalışan, 5 kademe
    python main.py --calisanlar calisanlar.xlsx --bantlar bantlar.xlsx --yil 2026
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import statistics
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
PARAMETRE = BURASI / "tr_parametreler.json"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
K2 = Decimal("0.01")
PERF = {"a": 4, "b": 3, "c": 2, "d": 1, "5": 5, "4": 4, "3": 3, "2": 2, "1": 1}

CALISAN_SUTUNLARI = {"sicil": ("sicil no", "sicil"), "ad": ("ad soyad", "ad"), "departman": ("departman", "birim"), "pozisyon": ("pozisyon", "unvan"),
                     "kademe": ("kademe", "grade", "seviye"), "ucret": ("brut ucret tl", "brut ucret", "ucret", "aylik brut"),
                     "giris": ("ise giris", "ise giris tarihi"), "cinsiyet": ("cinsiyet",), "performans": ("performans", "performans notu"),
                     "oran": ("calisma orani", "fte", "tam zaman orani")}
BANT_SUTUNLARI = {"kademe": ("kademe", "grade"), "pozisyon": ("pozisyon",), "alt": ("alt tl", "alt", "minimum", "p25"),
                  "orta": ("orta tl", "orta", "orta nokta", "medyan", "p50"), "ust": ("ust tl", "ust", "maksimum", "p75"), "kaynak": ("kaynak",)}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tl(x: Decimal | None) -> str:
    return "—" if x is None else f"{x.quantize(K2, ROUND_HALF_UP):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def yuzde(x: float) -> str:
    return f"%{x * 100:.1f}".replace(".", ",")


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


def asgari_ucret(yil: int) -> Decimal | None:
    if not PARAMETRE.exists():
        return None
    v = json.loads(PARAMETRE.read_text(encoding="utf-8")).get("yillar", {}).get(str(yil), {}).get("asgari_ucret_brut")
    return Decimal(v) if v else None


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Bant:
    kademe: str
    pozisyon: str
    alt: Decimal
    orta: Decimal
    ust: Decimal
    kaynak: str

    @property
    def ad(self) -> str:
        return f"{self.kademe}" + (f" / {self.pozisyon}" if self.pozisyon else "")


@dataclass
class Calisan:
    sicil: str
    ad: str
    departman: str
    pozisyon: str
    kademe: str
    ucret: Decimal
    giris: date | None
    cinsiyet: str
    performans: str
    oran: Decimal
    bant: Bant | None = None
    durum: str = ""
    notlar: list = field(default_factory=list)

    @property
    def tam_zaman(self) -> Decimal:
        return (self.ucret / self.oran).quantize(K2, ROUND_HALF_UP) if self.oran else self.ucret

    @property
    def compa(self) -> float | None:
        return float(self.tam_zaman / self.bant.orta) if self.bant else None

    @property
    def konum(self) -> float | None:
        if not self.bant or self.bant.ust == self.bant.alt:
            return None
        return float((self.tam_zaman - self.bant.alt) / (self.bant.ust - self.bant.alt))

    @property
    def perf(self) -> int | None:
        return PERF.get(katla(self.performans)[:1]) if self.performans else None


def oku_bantlar(yol: Path) -> list[Bant]:
    sonuc = []
    for r in kayitlar(yol, BANT_SUTUNLARI, ("alt", "ust")):
        alt, ust = para(r.get("alt")), para(r.get("ust"))
        if alt is None or ust is None:
            continue
        orta = para(r.get("orta")) or (alt + ust) / 2
        sonuc.append(Bant(metin(r.get("kademe")), metin(r.get("pozisyon")), alt, orta, ust, metin(r.get("kaynak"))))
    return sonuc


def oku_calisanlar(yol: Path) -> tuple[list[Calisan], list[str]]:
    sonuc, hatalar = [], []
    for r in kayitlar(yol, CALISAN_SUTUNLARI, ("sicil", "ucret")):
        if not r.get("sicil"):
            continue
        u = para(r.get("ucret"))
        if u is None:
            hatalar.append(f"Satır {r['_satir']} ({r['sicil']}): ücret okunamadı")
            continue
        o = para(r.get("oran")) or Decimal(1)
        sonuc.append(Calisan(metin(r["sicil"]), metin(r.get("ad")), metin(r.get("departman")) or "—", metin(r.get("pozisyon")), metin(r.get("kademe")), u,
                             tarih(r.get("giris")), metin(r.get("cinsiyet")), metin(r.get("performans")), o / 100 if o > 1 else o))
    return sonuc, hatalar


# ----------------------------------------------------------------------------
# Analiz
# ----------------------------------------------------------------------------

SIRA = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}


def analiz_et(calisanlar: list[Calisan], bantlar: list[Bant], asgari: Decimal | None, sikisma_orani: float = 0.95, sikisma_yil: int = 2,
              cinsiyet_esik: float = 0.05) -> dict:
    uyarilar = []

    def uyar(onem, tur, c: Calisan | None, aciklama):
        uyarilar.append({"onem": onem, "tur": tur, "sicil": c.sicil if c else "", "ad": c.ad if c else "", "aciklama": aciklama})
        if c:
            c.notlar.append(tur)

    for c in calisanlar:
        c.bant = next((b for b in bantlar if b.pozisyon and katla(b.pozisyon) == katla(c.pozisyon)), None) or \
            next((b for b in bantlar if not b.pozisyon and katla(b.kademe) == katla(c.kademe)), None)
        u = c.tam_zaman
        if asgari is not None and u < asgari:
            uyar("Yüksek", "Asgari ücret altı", c, f"Tam zamanlı karşılık {tl(u)} TL < yasal asgari ücret {tl(asgari)} TL")
        if c.bant is None:
            c.durum = "Bant yok"
            uyar("Bilgi", "Bant tanımsız", c, f"Kademe '{c.kademe}' / pozisyon '{c.pozisyon}' için bant yok")
            continue
        if u < c.bant.alt:
            c.durum = "Bant altı"
            uyar("Yüksek", "Bant altı", c, f"{tl(u)} TL < bant altı {tl(c.bant.alt)} TL ({c.bant.ad}); fark {tl(c.bant.alt - u)} TL / ay"
                 + (f" (kısmi süreli, oran {c.oran:g})" if c.oran != 1 else ""))
        elif u > c.bant.ust:
            c.durum = "Bant üstü"
            uyar("Orta", "Bant üstü", c, f"{tl(u)} TL > bant üstü {tl(c.bant.ust)} TL ({c.bant.ad}); artışlar tek seferlik ödeme ile "
                 "değerlendirilebilir veya kademe gözden geçirilmeli")
        else:
            c.durum = "Bant içi"
    # Ücret sıkışması
    gruplar = defaultdict(list)
    for c in calisanlar:
        gruplar[katla(c.pozisyon)].append(c)
    for lst in gruplar.values():
        for a in lst:
            if not a.giris:
                continue
            yeniler = [b for b in lst if b is not a and b.giris and (b.giris - a.giris).days >= 365 * sikisma_yil
                       and (a.perf is None or b.perf is None or a.perf >= b.perf) and float(a.tam_zaman) < float(b.tam_zaman) * sikisma_orani]
            if yeniler:
                b = max(yeniler, key=lambda x: x.tam_zaman)
                uyar("Orta", "Ücret sıkışması", a, f"{a.giris:%Y} girişli, performans {a.performans or '—'}, ücret {tl(a.tam_zaman)} TL; "
                     f"{b.giris:%Y} girişli {b.sicil} (performans {b.performans or '—'}) " + ("tam zamanlı karşılık " if b.oran != 1 else "")
                     + f"{tl(b.tam_zaman)} TL alıyor")
    # Bütçe etkisi
    maliyet = [(c, (c.bant.alt - c.tam_zaman) * c.oran) for c in calisanlar if c.durum == "Bant altı"]
    # Özetler
    kademe = []
    g = defaultdict(list)
    for c in calisanlar:
        if c.bant:
            g[c.bant.ad].append(c)
    for ad, lst in sorted(g.items()):
        cr = [c.compa for c in lst]
        kademe.append({"bant": lst[0].bant, "n": len(lst), "ort_compa": statistics.fmean(cr), "min": min(cr), "max": max(cr),
                       "alt": sum(c.durum == "Bant altı" for c in lst), "ust": sum(c.durum == "Bant üstü" for c in lst),
                       "medyan_ucret": statistics.median(float(c.tam_zaman) for c in lst)})
    departman = []
    g = defaultdict(list)
    for c in calisanlar:
        g[c.departman].append(c)
    for ad, lst in sorted(g.items()):
        cr = [c.compa for c in lst if c.compa is not None]
        departman.append({"ad": ad, "n": len(lst), "ort_compa": statistics.fmean(cr) if cr else None, "alt": sum(c.durum == "Bant altı" for c in lst),
                          "ust": sum(c.durum == "Bant üstü" for c in lst), "toplam": sum((c.ucret for c in lst), Decimal(0))})
    cinsiyet = []
    g = defaultdict(lambda: defaultdict(list))
    for c in calisanlar:
        if c.bant and c.cinsiyet:
            g[c.bant.ad][c.cinsiyet].append(c.compa)
    for ad, grup in sorted(g.items()):
        if len(grup) >= 2 and all(len(v) >= 2 for v in grup.values()):
            ort = {k: statistics.fmean(v) for k, v in grup.items()}
            fark = (max(ort.values()) - min(ort.values())) / max(ort.values())
            cinsiyet.append({"bant": ad, "ort": ort, "n": {k: len(v) for k, v in grup.items()}, "fark": fark})
            if fark > cinsiyet_esik:
                uyar("Bilgi", "Cinsiyete göre fark", None, f"{ad}: ortalama karşılaştırma oranı " + ", ".join(f"{k} {v:.2f}" for k, v in ort.items())
                     + f" (fark {yuzde(fark)}). Kıdem, performans ve iş içeriği farklarıyla açıklanıp açıklanmadığı incelenmeli.")
    uyarilar.sort(key=lambda u: (SIRA[u["onem"]], u["tur"], u["sicil"]))
    return {"uyarilar": uyarilar, "maliyet": maliyet, "kademe": kademe, "departman": departman, "cinsiyet": cinsiyet}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Bant altı": "FDE2E1", "Bant üstü": "FFF4CE", "Bant içi": "E3F4E1", "Bant yok": "EEEEEE", "Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KONTROL = PatternFill("solid", fgColor="FFF4CE")
UST = Alignment(vertical="top", wrap_text=True)
PF = "#,##0.00"


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def rapor_yaz(cikti: Path, calisanlar: list[Calisan], s: dict, asgari: Decimal | None, yil: int) -> None:
    wb = Workbook()
    cl = wb.active
    cl.title = "Çalışanlar"
    _baslik(cl, ["Sicil", "Ad Soyad", "Departman", "Pozisyon", "Kademe", "Brüt Ücret", "Çalışma Oranı", "Tam Zamanlı Karşılık", "Bant", "Alt", "Orta", "Üst",
                 "Karşılaştırma Oranı", "Bant İçi Konum", "Durum", "İşaretler", "Öneri / Karar"],
            (8, 16, 14, 22, 7, 12, 8, 13, 22, 11, 11, 11, 11, 10, 11, 26, 24))
    for c in sorted(calisanlar, key=lambda c: (c.departman, c.kademe, c.sicil)):
        b = c.bant
        cl.append([c.sicil, c.ad, c.departman, c.pozisyon, c.kademe, float(c.ucret), float(c.oran), float(c.tam_zaman), b.ad if b else "—",
                   b and float(b.alt), b and float(b.orta), b and float(b.ust), c.compa, c.konum, c.durum, "; ".join(c.notlar), ""])
        r = cl.max_row
        for j in (6, 8, 10, 11, 12):
            cl.cell(r, j).number_format = PF
        cl.cell(r, 13).number_format = "0.00"
        cl.cell(r, 14).number_format = "0%"
        cl.cell(r, 15).fill = PatternFill("solid", fgColor=RENK[c.durum])
        cl.cell(r, 17).fill = KONTROL
    cl.auto_filter.ref = f"A1:Q{cl.max_row}"

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Sicil", "Ad Soyad", "Açıklama", "Karar"], (9, 20, 8, 16, 90, 24))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["sicil"], u["ad"], u["aciklama"], ""])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 5).alignment = UST
        uy.cell(uy.max_row, 6).fill = KONTROL

    bt = wb.create_sheet("Bütçe Etkisi")
    _baslik(bt, ["Sicil", "Ad Soyad", "Bant", "Mevcut Brüt", "Bant Altı", "Aylık Ek Maliyet (brüt)", "Yıllık Ek Maliyet (brüt)"], (8, 16, 22, 12, 12, 16, 16))
    toplam = Decimal(0)
    for c, m in s["maliyet"]:
        bt.append([c.sicil, c.ad, c.bant.ad, float(c.ucret), float(c.bant.alt), float(m), float(m * 12)])
        toplam += m
        for j in (4, 5, 6, 7):
            bt.cell(bt.max_row, j).number_format = PF
    bt.append(["Toplam", "", "", None, None, float(toplam), float(toplam * 12)])
    for j in (6, 7):
        bt.cell(bt.max_row, j).number_format = PF
        bt.cell(bt.max_row, j).font = Font(bold=True)
    bt.append(["Not: Brüt ücret üzerinden; işveren SGK ve işsizlik primi payı dahil değildir."])

    oz = wb.create_sheet("Kademe Özeti")
    _baslik(oz, ["Bant", "Kaynak", "Çalışan", "Alt", "Orta", "Üst", "Medyan Ücret", "Ort. Karşılaştırma Oranı", "En Düşük", "En Yüksek", "Bant Altı", "Bant Üstü"],
            (22, 34, 8, 11, 11, 11, 12, 12, 9, 9, 9, 9))
    for x in s["kademe"]:
        b = x["bant"]
        oz.append([b.ad, b.kaynak, x["n"], float(b.alt), float(b.orta), float(b.ust), x["medyan_ucret"], x["ort_compa"], x["min"], x["max"], x["alt"], x["ust"]])
        for j in (4, 5, 6, 7):
            oz.cell(oz.max_row, j).number_format = PF
        for j in (8, 9, 10):
            oz.cell(oz.max_row, j).number_format = "0.00"
    oz.append([])
    _baslik(oz, ["Departman", "Çalışan", "Ort. Karşılaştırma Oranı", "Bant Altı", "Bant Üstü", "Aylık Brüt Toplam"], ())
    for x in s["departman"]:
        oz.append([x["ad"], x["n"], x["ort_compa"], x["alt"], x["ust"], float(x["toplam"])])
        oz.cell(oz.max_row, 3).number_format = "0.00"
        oz.cell(oz.max_row, 6).number_format = PF
    oz.append([])
    oz.append([f"Asgari ücret ({yil}, brüt): {tl(asgari)} TL" if asgari else f"Asgari ücret: {yil} için tr_parametreler.json'da değer yok"])

    cs = wb.create_sheet("Cinsiyet Karşılaştırması")
    _baslik(cs, ["Bant", "Grup", "Kişi", "Ort. Karşılaştırma Oranı", "Gruplar Arası Fark"], (22, 12, 8, 14, 14))
    for x in s["cinsiyet"]:
        for k, v in x["ort"].items():
            cs.append([x["bant"], k, x["n"][k], v, x["fark"]])
            cs.cell(cs.max_row, 4).number_format = "0.00"
            cs.cell(cs.max_row, 5).number_format = "0.0%"
    cs.append([])
    cs.append(["Her grupta en az 2 kişi olan bantlar gösterilir. Fark tek başına ayrımcılık göstermez; kıdem, performans ve iş içeriğiyle birlikte "
               "değerlendirilmelidir (4857 sayılı Kanun md. 5)."])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(calisan_yolu: Path, bant_yolu: Path, cikti: Path, yil: int | None = None, sikisma_orani: float = 0.95) -> dict:
    yil = yil or date.today().year
    calisanlar, hatalar = oku_calisanlar(calisan_yolu)
    asgari = asgari_ucret(yil)
    s = analiz_et(calisanlar, oku_bantlar(bant_yolu), asgari, sikisma_orani)
    for h in hatalar:
        s["uyarilar"].append({"onem": "Orta", "tur": "Okunamayan satır", "sicil": "", "ad": "", "aciklama": h})
    rapor_yaz(cikti, calisanlar, s, asgari, yil)
    return {**s, "calisanlar": calisanlar, "asgari": asgari}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="İç ücretleri kademe / pozisyon bazında ücret araştırması bantlarıyla kıyaslar, bant dışı çalışanları işaretler.")
    p.add_argument("--calisanlar", type=Path, default=ORNEK / "calisanlar.csv", help="Sicil, Ad, Departman, Pozisyon, Kademe, Brüt Ücret, İşe Giriş, …")
    p.add_argument("--bantlar", type=Path, default=ORNEK / "bantlar.csv", help="Kademe, Pozisyon (isteğe bağlı), Alt, Orta, Üst, Kaynak")
    p.add_argument("--yil", type=int, help="Asgari ücret yılı (varsayılan bu yıl; örnek veride 2026)")
    p.add_argument("--sikisma", type=float, default=95, help="Kıdemli çalışan yeni gelenin ücretinin bu %%'sinden azsa sıkışma (varsayılan 95)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "ucret_bandi_karsilastirmasi.xlsx")
    a = p.parse_args(argv)
    for y in (a.calisanlar, a.bantlar):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    yil = a.yil or (2026 if a.calisanlar == ORNEK / "calisanlar.csv" else None)
    try:
        s = calistir(a.calisanlar, a.bantlar, a.cikti, yil, a.sikisma / 100)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    durum = defaultdict(int)
    for c in s["calisanlar"]:
        durum[c.durum] += 1
    print(f"[OK] {len(s['calisanlar'])} çalışan · " + " · ".join(f"{k} {v}" for k, v in durum.items()))
    for u in s["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['sicil']} {u['tur']}: {u['aciklama']}")
    m = sum((x for _, x in s["maliyet"]), Decimal(0))
    print(f"[OK] Bant altını alt sınıra çekme maliyeti: aylık {tl(m)} TL, yıllık {tl(m * 12)} TL (brüt) · Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
