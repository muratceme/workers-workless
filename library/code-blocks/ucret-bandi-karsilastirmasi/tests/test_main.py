import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "u.xlsx"
        cls.s = main.calistir(ORNEK / "calisanlar.csv", ORNEK / "bantlar.csv", cls.cikti, 2026)
        cls.c = {c.sicil: c for c in cls.s["calisanlar"]}
        cls.u = {}
        for u in cls.s["uyarilar"]:
            cls.u.setdefault(u["tur"], []).append(u["sicil"])

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_bant_ve_oranlar(self):
        self.assertEqual(self.s["asgari"], D("33030.00"))
        self.assertEqual(self.c["1017"].bant.ad, "K3 / Yazılım Geliştirici")
        self.assertEqual(self.c["1012"].bant.ad, "K3")
        self.assertEqual(self.c["1002"].tam_zaman, D("41333.33"))
        self.assertEqual(self.c["1002"].durum, "Bant içi")
        self.assertAlmostEqual(self.c["1021"].compa, 59000 / 48000)
        self.assertAlmostEqual(self.c["1009"].konum, (40000 - 42000) / (56000 - 42000))

    def test_uyarilar(self):
        self.assertEqual(self.u["Asgari ücret altı"], ["1001"])
        self.assertEqual(sorted(self.u["Bant altı"]), ["1001", "1009", "1017"])
        self.assertEqual(self.u["Bant üstü"], ["1021"])
        self.assertEqual(sorted(self.u["Ücret sıkışması"]), ["1003", "1008", "1009", "1019"])
        self.assertNotIn("Cinsiyete göre fark", self.u)

    def test_maliyet_ve_ozet(self):
        self.assertEqual(sum(m for _, m in self.s["maliyet"]), D("9500.00"))
        k = {x["bant"].ad: x for x in self.s["kademe"]}
        self.assertEqual((k["K1"]["n"], k["K1"]["alt"]), (8, 1))
        self.assertEqual(len(self.s["cinsiyet"]), 3)

    def test_excel(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Çalışanlar", "Uyarılar", "Bütçe Etkisi", "Kademe Özeti", "Cinsiyet Karşılaştırması"])
        self.assertEqual(wb["Bütçe Etkisi"]["G5"].value, 114000)
        self.assertEqual(wb["Çalışanlar"]["Q1"].value, "Öneri / Karar")


class KuralTesti(unittest.TestCase):
    def test_cinsiyet_farki_ve_bant_yok(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t / "b.csv").write_text("Kademe;Alt;Orta;Üst\nK1;100;120;140\n", encoding="utf-8")
            (t / "c.csv").write_text("Sicil No;Kademe;Brüt Ücret (TL);Cinsiyet\n1;K1;100;Kadın\n2;K1;104;Kadın\n3;K1;130;Erkek\n4;K1;134;Erkek\n5;K9;50;Erkek\n",
                                     encoding="utf-8")
            s = main.calistir(t / "c.csv", t / "b.csv", t / "o.xlsx", 2026)
            turler = {u["tur"] for u in s["uyarilar"]}
            self.assertIn("Cinsiyete göre fark", turler)
            self.assertIn("Bant tanımsız", turler)
            self.assertIn("Asgari ücret altı", turler)

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--bantlar", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
