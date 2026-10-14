"""
Periyodik Bakım Hatırlatma Listesi — Workers / Workless kod bloğu
Satış Sonrası Hizmetler (Servis) › Servis Danışmanı
Son bakım tarihi ve kilometre tahminine göre bakımı yaklaşan ya da geçen araçları listeler, hatırlatma mesajı hazırlar.
  Günlük km       = (son bakım km − önceki bakım km) ÷ iki bakım arası gün; önceki bakım yoksa --gunluk-km (40)
  Tahmini km      = son bakım km + günlük km × (bugün − son bakım tarihi)
  Km'ye göre vade = son bakım tarihi + ⌈bakım aralığı km ÷ günlük km⌉ gün
  Süreye göre vade= son bakım tarihi + bakım aralığı ay
  Bakım vadesi    = hangisi önce gelirse; kalan gün = vade − bugün
  Liste           = vadesi geçmiş ya da --pencere (30) gün içinde olanlar; mesaj yalnız ileti izni "Evet" olanlara
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
SUTUN = {"plaka": ("plaka",), "musteri": ("musteri", "ad soyad", "musteri adi"), "telefon": ("telefon", "cep telefonu", "gsm", "tel"), "model": ("model", "arac", "marka model"),
         "son_tarih": ("son bakim tarihi", "son bakim", "son servis tarihi"), "son_km": ("son bakim km", "son km", "son servis km"),
         "onceki_tarih": ("onceki bakim tarihi", "onceki bakim", "onceki servis tarihi"), "onceki_km": ("onceki bakim km", "onceki km", "onceki servis km"),
         "aralik_km": ("bakim araligi km", "aralik km", "bakim km"), "aralik_ay": ("bakim araligi ay", "aralik ay", "bakim ay"),
         "izin": ("ileti izni", "iys izni", "izin", "ticari ileti onayi", "iletisim izni")}
MESAJ = "Sayın {musteri}, {plaka} plakalı aracınızın periyodik bakım zamanı {tarih}. Randevu için bizi arayabilirsiniz. {firma}"


def ay_ekle(g: date, ay: int) -> date:
    y, m = divmod(g.month - 1 + ay, 12)
    yil, ayn = g.year + y, m + 1
    son = (date(yil + (ayn == 12), ayn % 12 + 1, 1) - timedelta(days=1)).day
    return date(yil, ayn, min(g.day, son))


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("plaka", "son_tarih")):
        plaka = " ".join(metin(r["plaka"]).upper().split())
        if not plaka:
            continue
        st = tarih(r["son_tarih"])
        sayi = {a: (para(r[a]) if metin(r[a]) else None) for a in ("son_km", "onceki_km", "aralik_km", "aralik_ay")}
        ot = tarih(r["onceki_tarih"]) if metin(r["onceki_tarih"]) else None
        if st is None or (metin(r["onceki_tarih"]) and ot is None) or any(metin(r[a]) and (v is None or v < 0) for a, v in sayi.items()):
            uy.append(uyari("Yüksek", "Satır okunamadı", plaka, f"Satır {r['_satir']}: son bakım tarihi, önceki bakım tarihi ya da km / aralık değerlerinden biri okunamadı."))
            continue
        if plaka in gorulen:
            uy.append(uyari("Orta", "Plaka iki kez yazılmış", plaka, f"Satır {r['_satir']}: ikinci satır alınmadı."))
            continue
        gorulen.add(plaka)
        l.append({"plaka": plaka, "musteri": metin(r["musteri"]), "telefon": metin(r["telefon"]), "model": metin(r["model"]), "son_tarih": st, "onceki_tarih": ot,
                  **sayi, "izin": katla(r["izin"]) in ("evet", "e", "var", "onayli", "x", "1")})
    return l, uy


def hesapla(araclar: list[dict], bugun: date, pencere: int = 30, gunluk_km: Decimal = Decimal(40), aralik_km: Decimal | None = Decimal(15000), aralik_ay: int | None = 12,
            mesaj: str = MESAJ, firma: str = "") -> dict:
    if not araclar:
        raise ValueError("Okunabilen araç yok.")
    uy, varsayilan = [], 0
    for a in araclar:
        a["km_kaynak"] = "Varsayılan"
        a["gunluk_km"] = gunluk_km
        if a["onceki_tarih"] and a["onceki_km"] is not None and a["son_km"] is not None:
            gun = (a["son_tarih"] - a["onceki_tarih"]).days
            fark = a["son_km"] - a["onceki_km"]
            if gun > 0 and fark >= 0:
                a["gunluk_km"], a["km_kaynak"] = yuvarla(fark / gun, "0.1"), "Geçmiş bakım"
            else:
                uy.append(uyari("Orta", "Km / tarih tutarsız", a["plaka"], f"Önceki bakım {a['onceki_tarih']:%d.%m.%Y} · {a['onceki_km']} km, son bakım {a['son_tarih']:%d.%m.%Y} · {a['son_km']} km; varsayılan günlük km kullanıldı."))
        akm = a["aralik_km"] if a["aralik_km"] else aralik_km
        aay = int(a["aralik_ay"]) if a["aralik_ay"] else aralik_ay
        if not a["aralik_km"] or not a["aralik_ay"]:
            varsayilan += 1
        gecen = (bugun - a["son_tarih"]).days
        a["tahmini_km"] = int(a["son_km"] + a["gunluk_km"] * gecen) if a["son_km"] is not None and gecen >= 0 else None
        vadeler = []
        if akm and a["gunluk_km"] > 0:
            vadeler.append((a["son_tarih"] + timedelta(days=math.ceil(akm / a["gunluk_km"])), "Km"))
        if aay:
            vadeler.append((ay_ekle(a["son_tarih"], aay), "Süre"))
        if not vadeler:
            a.update({"vade": None, "neden": "", "kalan": None, "durum": "Hesaplanamadı", "hedef_km": None})
            continue
        a["vade"], a["neden"] = min(vadeler)
        a["kalan"] = (a["vade"] - bugun).days
        a["hedef_km"] = int(a["son_km"] + akm) if a["son_km"] is not None and akm else None
        a["durum"] = "Gecikmiş" if a["kalan"] < 0 else "Yaklaşıyor" if a["kalan"] <= pencere else "İleride"
        if a["kalan"] < -180:
            uy.append(uyari("Bilgi", "Uzun süredir gelmiyor", a["plaka"], f"Bakım vadesi {-a['kalan']} gün önce geçti; araç başka serviste bakım yaptırıyor olabilir (kayıp müşteri araması)."))
    if varsayilan:
        uy.append(uyari("Bilgi", "Varsayılan bakım aralığı", f"{varsayilan} araç", f"Bakım aralığı yazılmayan araçlarda {aralik_km} km / {aralik_ay} ay kullanıldı; modelin bakım planından doğrulayın."))
    liste = sorted((a for a in araclar if a["durum"] in ("Gecikmiş", "Yaklaşıyor")), key=lambda a: (a["kalan"], a["plaka"]))
    mesajlar = []
    for a in liste:
        if not a["izin"]:
            continue
        if not a["telefon"]:
            uy.append(uyari("Orta", "Telefon yok", a["plaka"], "İleti izni var ama telefon numarası yazılmamış."))
            continue
        tarih_metni = "geldi" if a["kalan"] <= 0 else f"{a['vade']:%d.%m.%Y} civarında geliyor"
        m = mesaj.format(musteri=a["musteri"] or "Müşterimiz", plaka=a["plaka"], tarih=tarih_metni, model=a["model"], firma=firma).strip()
        mesajlar.append({"plaka": a["plaka"], "telefon": a["telefon"], "mesaj": " ".join(m.split())})
    izinsiz = sum(1 for a in liste if not a["izin"])
    if izinsiz:
        uy.append(uyari("Bilgi", "İleti izni yok", f"{izinsiz} araç", "Listedeki bu araçların sahiplerine ticari elektronik ileti izni olmadan SMS / e-posta gönderilmez; mesaj hazırlanmadı."))
    ozet = {d: sum(1 for a in araclar if a["durum"] == d) for d in ("Gecikmiş", "Yaklaşıyor", "İleride", "Hesaplanamadı")}
    return {"araclar": sorted(araclar, key=lambda a: (a["kalan"] is None, a["kalan"] if a["kalan"] is not None else 0, a["plaka"])), "liste": liste, "mesajlar": mesajlar,
            "ozet": ozet, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Hatırlatma Listesi"
    bas = ["Plaka", "Müşteri", "Telefon", "Model", "Son Bakım", "Son Bakım Km", "Günlük Km", "Km Kaynağı", "Tahmini Km (bugün)", "Hedef Km", "Bakım Vadesi", "Vade Nedeni", "Kalan Gün",
           "Durum", "İleti İzni"]
    gen = (12, 20, 14, 18, 11, 10, 8, 12, 11, 10, 11, 8, 8, 11, 7)
    for ad, l in (("Hatırlatma Listesi", r["liste"]), ("Tüm Araçlar", r["araclar"])):
        if ad != "Hatırlatma Listesi":
            ws = wb.create_sheet(ad)
        _baslik(ws, bas, gen, "B2")
        for a in l:
            _satir(ws, [a["plaka"], a["musteri"], a["telefon"], a["model"], a["son_tarih"], a["son_km"], a["gunluk_km"], a["km_kaynak"], a["tahmini_km"], a["hedef_km"], a["vade"], a["neden"],
                        a["kalan"], a["durum"], "Evet" if a["izin"] else "Hayır"], tarih_s=(5, 11),
                   dolgu={14: KIRMIZI if a["durum"] == "Gecikmiş" else SARI if a["durum"] == "Yaklaşıyor" else None, 15: None if a["izin"] else MAVI})
    ws = wb.create_sheet("Mesajlar")
    _baslik(ws, ["Plaka", "Telefon", "Mesaj"], (12, 14, 110))
    for m in r["mesajlar"]:
        _satir(ws, [m["plaka"], m["telefon"], m["mesaj"]], sar=(3,))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, **k) -> dict:
    araclar, uy = oku(yol)
    r = hesapla(araclar, bugun, **k)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    with open(cikti.with_name(cikti.stem + "_mesajlar.csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Telefon", "Mesaj"])
        for m in r["mesajlar"]:
            w.writerow([m["telefon"], m["mesaj"]])
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Bakımı yaklaşan ve geçen araçları listeler, hatırlatma mesajı hazırlar.")
    p.add_argument("--araclar", type=Path, help="Plaka, Müşteri, Telefon, Model, Son Bakım Tarihi, Son Bakım Km, Önceki Bakım Tarihi, Önceki Bakım Km, Bakım Aralığı Km, Bakım Aralığı Ay, İleti İzni")
    p.add_argument("--tarih", help="Rapor günü GG.AA.YYYY (varsayılan bugün; örnek veride 14.10.2026)")
    p.add_argument("--pencere", type=int, default=30, help="Kaç gün içindeki bakımlar listelenir (varsayılan 30)")
    p.add_argument("--gunluk-km", default="40", help="Önceki bakımı olmayan araçlar için günlük km (varsayılan 40)")
    p.add_argument("--aralik-km", default="15000", help="Satırda yoksa bakım aralığı km (varsayılan 15000)")
    p.add_argument("--aralik-ay", type=int, default=12, help="Satırda yoksa bakım aralığı ay (varsayılan 12)")
    p.add_argument("--firma", default="", help="Mesaj sonuna eklenecek servis adı")
    p.add_argument("--mesaj", default=MESAJ, help="Mesaj şablonu: {musteri} {plaka} {tarih} {model} {firma}")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "bakim_hatirlatma.xlsx")
    a = p.parse_args(argv)
    yol = a.araclar or ORNEK / "araclar.csv"
    bugun = tarih(a.tarih) if a.tarih else (date(2026, 10, 14) if a.araclar is None else date.today())
    gkm, akm = para(a.gunluk_km), para(a.aralik_km)
    try:
        a.mesaj.format(musteri="", plaka="", tarih="", model="", firma="")
        sablon = True
    except (KeyError, IndexError, ValueError):
        sablon = False
    if bugun is None or gkm is None or gkm <= 0 or akm is None or akm <= 0 or a.aralik_ay < 1 or a.pencere < 0 or not sablon:
        print("[X] --tarih GG.AA.YYYY; --gunluk-km, --aralik-km, --aralik-ay 0'dan büyük; --pencere 0 ya da üstü; --mesaj yalnız {musteri} {plaka} {tarih} {model} {firma} içerebilir")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, bugun, pencere=a.pencere, gunluk_km=gkm, aralik_km=akm, aralik_ay=a.aralik_ay, mesaj=a.mesaj, firma=a.firma)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {len(r['araclar'])} araç · gecikmiş {o['Gecikmiş']} · {a.pencere} gün içinde {o['Yaklaşıyor']} · {len(r['mesajlar'])} mesaj hazırlandı")
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Liste: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
