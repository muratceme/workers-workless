"""
Ürün Gamı (Delist) Önerisi — Gerisi Bende kod bloğu
Satın Alma ve Kategori Yönetimi › Kategori Yöneticisi
Satış, marj ve raf verimliliğine göre listeden çıkarılabilecek ürünleri önerir. Her ürün kendi kategorisi içinde değerlendirilir.
  Ciro kuyruğu       = ciroya göre sıralandığında, kendisinden önceki ürünler kategori cirosunun %(100 − --kuyruk) kadarını
                       zaten oluşturuyorsa (varsayılan --kuyruk 10: cironun son %10'unu oluşturan ürünler)
  Kâr kuyruğu        = aynı kural brüt kâra göre; brüt kârı sıfır ya da eksi olan ürün her zaman kuyruktadır
  Raf verimliliği    = kâr payı ÷ raf payı × 100; --raf (50) altı düşük sayılır (raf alanı yazılmışsa)
  Öneri              = değerlendirilebilen ölçütlerin (en az iki) hepsi olumsuz → Delist adayı · ikisi olumsuz → İzle · diğer → Devam
  Hariç              = Stratejik işaretli ürünler · listeleme tarihi --yeni (90) günden yeni ürünler
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"kod": ("urun kodu", "stok kodu", "kod", "malzeme kodu"), "ad": ("urun adi", "urun", "aciklama"), "kategori": ("kategori", "urun grubu", "alt kategori", "reyon"),
         "adet": ("satis adedi", "satis miktari", "adet"), "ciro": ("ciro", "net satis", "satis tutari"), "kar": ("brut kar", "kar", "brut kar tutari"), "maliyet": ("maliyet", "satis maliyeti", "smm"),
         "raf": ("raf alani", "raf", "yuz sayisi", "facing", "raf cm"), "liste": ("listeleme tarihi", "giris tarihi", "acilis tarihi"), "stratejik": ("stratejik", "zorunlu", "korunacak")}
SIFIR = Decimal(0)
EVET = {"evet", "e", "x", "var", "1", "stratejik", "zorunlu"}


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("kod", "ciro")):
        kod = metin(r["kod"])
        if not kod:
            continue
        d = {a: para(r[a]) if metin(r[a]) else None for a in ("adet", "ciro", "kar", "maliyet", "raf")}
        liste = tarih(r["liste"]) if metin(r["liste"]) else None
        if d["ciro"] is None or any(metin(r[a]) and d[a] is None for a in d) or any(d[a] is not None and d[a] < 0 for a in ("adet", "ciro", "maliyet", "raf")) or (metin(r["liste"]) and liste is None):
            uy.append(uyari("Yüksek", "Satır okunamadı", kod, f"Satır {r['_satir']}: sayılardan ya da listeleme tarihinden biri okunamadı. Ürün alınmadı."))
            continue
        kar = d["kar"] if d["kar"] is not None else d["ciro"] - d["maliyet"] if d["maliyet"] is not None else None
        if kar is None:
            uy.append(uyari("Yüksek", "Kâr bilgisi yok", kod, f"Satır {r['_satir']}: Brüt Kâr ya da Maliyet sütunu gerekli. Ürün alınmadı."))
            continue
        kategori = metin(r["kategori"]) or "Diğer"
        if (katla(kategori), katla(kod)) in gorulen:
            uy.append(uyari("Yüksek", "Ürün iki kez yazılmış", kod, f"Satır {r['_satir']}: ikinci satır alınmadı."))
            continue
        gorulen.add((katla(kategori), katla(kod)))
        l.append({"kod": kod, "ad": metin(r["ad"]), "kategori": kategori, "adet": d["adet"], "ciro": d["ciro"], "kar": kar, "raf": d["raf"], "liste": liste, "stratejik": katla(r["stratejik"]) in EVET})
    return l, uy


def _pay(x, top) -> Decimal | None:
    return yuvarla(Decimal(x) * 100 / top) if top else None


def _kuyruk(urunler: list[dict], alan: str, kuyruk: Decimal) -> set[int]:
    top = sum((max(u[alan], SIFIR) for u in urunler), SIFIR)
    sinir, birikim, sonuc = top * (100 - kuyruk) / 100, SIFIR, set()
    for i, u in sorted(enumerate(urunler), key=lambda x: (-x[1][alan], katla(x[1]["kod"]))):
        if u[alan] <= 0 or birikim >= sinir:
            sonuc.add(i)
        birikim += max(u[alan], SIFIR)
    return sonuc


def degerlendir(urunler: list[dict], gun: date, kuyruk: Decimal = Decimal(10), raf_esik: Decimal = Decimal(50), yeni: int = 90) -> dict:
    if not urunler:
        raise ValueError("Değerlendirilecek ürün yok.")
    gruplar, uy, sonuc, ozet = {}, [], [], []
    for u in urunler:
        gruplar.setdefault(katla(u["kategori"]), []).append(u)
    for k in sorted(gruplar):
        l = gruplar[k]
        ciro, kar = sum((u["ciro"] for u in l), SIFIR), sum((u["kar"] for u in l), SIFIR)
        rafli = [u for u in l if u["raf"]]
        raf = sum((u["raf"] for u in rafli), SIFIR)
        kar_rafli = sum((max(u["kar"], SIFIR) for u in rafli), SIFIR)
        ck, kk = _kuyruk(l, "ciro", kuyruk), _kuyruk(l, "kar", kuyruk)
        if len(l) == 1:
            uy.append(uyari("Bilgi", "Kategoride tek ürün", l[0]["kategori"], "Karşılaştırılacak başka ürün olmadığı için öneri anlamlı değildir."))
        satirlar = []
        for i, u in enumerate(l):
            s = {**u, "ciro_pay": _pay(u["ciro"], ciro), "kar_pay": _pay(u["kar"], kar) if kar > 0 else None, "marj": _pay(u["kar"], u["ciro"]), "raf_pay": _pay(u["raf"], raf) if u["raf"] else None,
                 "ciro_kuyruk": i in ck, "kar_kuyruk": i in kk}
            s["raf_verim"] = yuvarla(max(u["kar"], SIFIR) * raf * 100 / (kar_rafli * u["raf"])) if u["raf"] and kar_rafli else (SIFIR if u["raf"] else None)
            s["raf_dusuk"] = None if s["raf_verim"] is None else s["raf_verim"] < raf_esik
            olcut = [s["ciro_kuyruk"], s["kar_kuyruk"]] + ([s["raf_dusuk"]] if s["raf_dusuk"] is not None else [])
            s["puan"], s["olcut"] = sum(olcut), len(olcut)
            gerekce = [ad for ad, v in (("ciro kuyruğu", s["ciro_kuyruk"]), ("kâr kuyruğu", s["kar_kuyruk"]), ("raf verimi düşük", s["raf_dusuk"])) if v]
            if u["kar"] < 0:
                gerekce.append("zarar ediyor")
            if not u["adet"] and u["adet"] is not None:
                gerekce.append("satış yok")
            yas = (gun - u["liste"]).days if u["liste"] else None
            if u["stratejik"]:
                s["oneri"] = "Stratejik"
            elif yas is not None and yas < yeni:
                s["oneri"] = "Yeni ürün"
                gerekce.append(f"{yas} gündür listede")
            else:
                s["oneri"] = "Delist adayı" if s["puan"] == s["olcut"] else "İzle" if s["puan"] >= 2 else "Devam"
            s["gerekce"] = ", ".join(gerekce)
            satirlar.append(s)
        satirlar.sort(key=lambda s: (-s["ciro"], katla(s["kod"])))
        aday = [s for s in satirlar if s["oneri"] == "Delist adayı"]
        ozet.append({"kategori": l[0]["kategori"], "urun": len(l), "ciro": ciro, "kar": kar, "marj": _pay(kar, ciro), "raf": raf, "aday": len(aday), "izle": sum(1 for s in satirlar if s["oneri"] == "İzle"),
                     "aday_ciro": _pay(sum((s["ciro"] for s in aday), SIFIR), ciro), "aday_kar": _pay(sum((s["kar"] for s in aday), SIFIR), kar) if kar > 0 else None,
                     "aday_raf": _pay(sum((s["raf"] for s in aday if s["raf"]), SIFIR), raf) if raf else None})
        sonuc += satirlar
    return {"urunler": sonuc, "kategoriler": ozet, "uyarilar": uy}


ONERI_RENK = {"Delist adayı": KIRMIZI, "İzle": SARI, "Yeni ürün": MAVI, "Stratejik": MAVI, "Devam": YESIL}


def rapor_yaz(cikti: Path, r: dict, kuyruk: Decimal, raf_esik: Decimal) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Delist Adayları"
    bas = ["Kategori", "Ürün Kodu", "Ürün Adı", "Öneri", "Gerekçe", "Satış Adedi", "Ciro", "Ciro Payı %", "Brüt Kâr", "Kâr Payı %", "Marj %", "Raf Alanı", "Raf Payı %", "Raf Verimliliği", "Listeleme"]
    gen = (16, 12, 28, 13, 34, 10, 14, 9, 14, 9, 8, 9, 9, 10, 11)

    def yaz(ws, s):
        _satir(ws, [s["kategori"], s["kod"], s["ad"], s["oneri"], s["gerekce"], s["adet"], s["ciro"], s["ciro_pay"], s["kar"], s["kar_pay"], s["marj"], s["raf"], s["raf_pay"], s["raf_verim"], s["liste"]],
               para_s=(7, 8, 9, 10, 11, 13, 14), tarih_s=(15,), dolgu={4: ONERI_RENK[s["oneri"]]}, sar=(5,))
    _baslik(ws, bas, gen, "D2")
    for s in r["urunler"]:
        if s["oneri"] in ("Delist adayı", "İzle"):
            yaz(ws, s)
    ws.append([])
    ws.append([f"Kuyruk: kategori ciro / kârının son %{kuyruk}'u. Raf verimliliği = kâr payı ÷ raf payı × 100; {raf_esik} altı düşük. Öneri karar değildir; ikame ürünleri ve tedarikçi anlaşmalarını kontrol edin."])
    ws = wb.create_sheet("Tüm Ürünler")
    _baslik(ws, bas, gen, "D2")
    for s in r["urunler"]:
        yaz(ws, s)
    ws = wb.create_sheet("Kategoriler")
    _baslik(ws, ["Kategori", "Ürün Sayısı", "Ciro", "Brüt Kâr", "Marj %", "Raf Alanı", "Delist Adayı", "İzle", "Adayların Ciro Payı %", "Adayların Kâr Payı %", "Adayların Raf Payı %"], (18, 9, 14, 14, 8, 10, 9, 8, 12, 12, 12))
    for o in r["kategoriler"]:
        _satir(ws, [o["kategori"], o["urun"], o["ciro"], o["kar"], o["marj"], o["raf"] or None, o["aday"], o["izle"], o["aday_ciro"], o["aday_kar"], o["aday_raf"]], para_s=(3, 4, 5, 9, 10, 11))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, gun: date, kuyruk: Decimal = Decimal(10), raf_esik: Decimal = Decimal(50), yeni: int = 90) -> dict:
    urunler, uy = oku(yol)
    r = degerlendir(urunler, gun, kuyruk, raf_esik, yeni)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, kuyruk, raf_esik)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Satış, marj ve raf verimliliğine göre listeden çıkarılabilecek ürünleri önerir.")
    p.add_argument("--urunler", type=Path, help="Ürün Kodu, Ürün Adı, Kategori, Satış Adedi, Ciro, Brüt Kâr (ya da Maliyet), Raf Alanı, Listeleme Tarihi, Stratejik")
    p.add_argument("--tarih", help="Değerlendirme tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 30.09.2026)")
    p.add_argument("--kuyruk", default="10", help="Kategori ciro / kârının son yüzde kaçı kuyruk sayılır (varsayılan 10)")
    p.add_argument("--raf", default="50", help="Raf verimliliği bu değerin altındaysa düşük sayılır (varsayılan 50)")
    p.add_argument("--yeni", type=int, default=90, help="Bu kadar günden yeni listelenen ürünler değerlendirme dışı (varsayılan 90)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "urun_gami_delist_onerisi.xlsx")
    a = p.parse_args(argv)
    yol = a.urunler or ORNEK / "kategori_urun_performansi.csv"
    gun = tarih(a.tarih) if a.tarih else (date(2026, 9, 30) if a.urunler is None else date.today())
    kuyruk, raf = para(a.kuyruk), para(a.raf)
    if gun is None or kuyruk is None or not 0 < kuyruk < 100 or raf is None or raf < 0 or a.yeni < 0:
        print("[X] --tarih GG.AA.YYYY, --kuyruk 0-100 arası, --raf ve --yeni sıfır ya da üstü olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, gun, kuyruk, raf, a.yeni)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(r['urunler'])} ürün · {len(r['kategoriler'])} kategori · {sum(o['aday'] for o in r['kategoriler'])} delist adayı · {sum(o['izle'] for o in r['kategoriler'])} izlenecek")
    for s in r["urunler"]:
        if s["oneri"] == "Delist adayı":
            print(f"[!] {s['kategori']} · {s['kod']} {s['ad']}: {s['gerekce']}")
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
