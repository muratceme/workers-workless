import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
GUN = date(2026, 9, 30)


def u(kod, kat, ciro, kar, raf=None, adet=None, liste=None, stratejik=False):
    return {"kod": kod, "ad": "", "kategori": kat, "adet": None if adet is None else D(adet), "ciro": D(ciro), "kar": D(kar), "raf": None if raf is None else D(raf), "liste": liste, "stratejik": stratejik}


def veri(raf=True):
    r = (lambda x: x) if raf else (lambda x: None)
    return [u("A", "K", 600, 150, r(30)), u("B", "K", 300, 90, r(30)), u("C", "K", 80, 20, r(20)), u("D", "K", 20, -5, r(20), 2), u("E", "k", 0, 0, r(10), 0, stratejik=True),
            u("F", "L", 100, 10, liste=date(2026, 9, 1))]


class DegerlendirTesti(unittest.TestCase):
    def test_oneriler(self):
        r = main.degerlendir(veri(), GUN)
        self.assertEqual([(s["kod"], s["oneri"], s["ciro_pay"], s["kar_pay"], s["marj"], s["raf_verim"], s["ciro_kuyruk"], s["kar_kuyruk"], s["gerekce"]) for s in r["urunler"]],
                         [("A", "Devam", D("60.00"), D("58.82"), D("25.00"), D("211.54"), False, False, ""), ("B", "Devam", D("30.00"), D("35.29"), D("30.00"), D("126.92"), False, False, ""),
                          ("C", "Delist adayı", D("8.00"), D("7.84"), D("25.00"), D("42.31"), True, True, "ciro kuyruğu, kâr kuyruğu, raf verimi düşük"),
                          ("D", "Delist adayı", D("2.00"), D("-1.96"), D("-25.00"), D("0.00"), True, True, "ciro kuyruğu, kâr kuyruğu, raf verimi düşük, zarar ediyor"),
                          ("E", "Stratejik", D("0.00"), D("0.00"), None, D("0.00"), True, True, "ciro kuyruğu, kâr kuyruğu, raf verimi düşük, satış yok"),
                          ("F", "Yeni ürün", D("100.00"), D("100.00"), D("10.00"), None, False, False, "29 gündür listede")])
        o = r["kategoriler"][0]
        self.assertEqual((o["urun"], o["ciro"], o["kar"], o["marj"], o["raf"], o["aday"], o["izle"], o["aday_ciro"], o["aday_kar"], o["aday_raf"]),
                         (5, D(1000), D(255), D("25.50"), D(110), 2, 0, D("10.00"), D("5.88"), D("36.36")))
        self.assertEqual([(x["tur"], x["kim"]) for x in r["uyarilar"]], [("Kategoride tek ürün", "L")])

    def test_esikler(self):
        o = lambda **k: [s["oneri"] for s in main.degerlendir(veri(), GUN, **k)["urunler"]]   # noqa: E731
        self.assertEqual(o(kuyruk=D(50)), ["Devam", "İzle", "Delist adayı", "Delist adayı", "Stratejik", "Yeni ürün"])
        self.assertEqual(o(raf_esik=D(0), yeni=29), ["Devam", "Devam", "İzle", "İzle", "Stratejik", "Devam"])
        self.assertEqual([s["oneri"] for s in main.degerlendir(veri(False), GUN)["urunler"]][:4], ["Devam", "Devam", "Delist adayı", "Delist adayı"])
        with self.assertRaises(ValueError):
            main.degerlendir([], GUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "kategori_urun_performansi.csv", cikti, GUN)
            self.assertEqual([(s["kod"], s["oneri"]) for s in r["urunler"] if s["oneri"] != "Devam"],
                             [("B05", "İzle"), ("B07", "Stratejik"), ("B06", "Delist adayı"), ("B09", "Yeni ürün"), ("B08", "Delist adayı"), ("S04", "Stratejik"), ("S05", "Delist adayı"), ("S06", "Delist adayı")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Delist Adayları", "Tüm Ürünler", "Kategoriler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "u.csv"
            y.write_text("Kod;Net Satış;Maliyet;Raf;Giriş Tarihi;Zorunlu\nA;1.000;600;10;01.01.2026;x\na;5;1;;;\nB;x;1;;;\nC;5;;;;\nD;5;1;-1;;\nE;5;1;;yok;\n;5;5;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["kod"], x["ciro"], x["kar"], x["raf"], x["liste"], x["stratejik"], x["kategori"]) for x in l], [x["tur"] for x in uy]),
                             ([("A", D(1000), D(400), D(10), date(2026, 1, 1), True, "Diğer")], ["Ürün iki kez yazılmış", "Satır okunamadı", "Kâr bilgisi yok", "Satır okunamadı", "Satır okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--urunler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "30.09.2026", "--kuyruk", "20", "--raf", "40", "--yeni", "0"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--kuyruk", "0"], ["--raf", "-1"], ["--yeni", "-1"])], [2] * 4)
            self.assertEqual(main.main(["--urunler", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Kod;Raf\nA;5\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
