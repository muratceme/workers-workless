import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
VERI = [("A", 40, ""), ("B", 30, "A"), ("C", 50, "A"), ("D", 40, "B"), ("E", 6, "B"), ("F", 25, "C"), ("G", 15, "C"), ("H", 20, "D,E"), ("I", 18, "F,G"), ("J", 30, "H,I")]


def ops(veri=VERI):
    return [{"op": k, "ad": "", "sure": D(str(s)), "oncul": [x for x in o.split(",") if x]} for k, s, o in veri]


def atama(r):
    return [[o["op"] for o in i["ops"]] for i in r["istasyonlar"]]


class DengelemeTesti(unittest.TestCase):
    def test_konum_agirligi(self):
        r = main.dengele(ops(), D(80))
        self.assertEqual({o["op"]: o["agirlik"] for o in r["ops"]}, {"A": D(274), "B": D(126), "C": D(138), "D": D(90), "E": D(56), "F": D(73), "G": D(63), "H": D(50), "I": D(48), "J": D(30)})
        self.assertEqual(atama(r), [["A", "B", "E"], ["C", "F"], ["D", "G", "H"], ["I", "J"]])
        self.assertEqual([(i["sure"], i["bos"], i["doluluk"]) for i in r["istasyonlar"]], [(D(76), D(4), D("95.00")), (D(75), D(5), D("93.75")), (D(75), D(5), D("93.75")), (D(48), D(32), D("60.00"))])
        self.assertEqual((r["toplam"], r["n"], r["teorik"], r["verim"], r["gecikme"], r["duzgunluk"], r["gerceklesen"], r["saatlik"], r["bos"]),
                         (D(274), 4, 4, D("85.63"), D("14.37"), D("28.04"), D(76), D("45.00"), D(46)))

    def test_oncelik_ihlali_olmaz(self):
        for c in (50, 60, 76, 80, 100, 274):
            for y in ("rpw", "lct"):
                r = main.dengele(ops(), D(c), y)
                ist = {o["op"]: o["istasyon"] for o in r["ops"]}
                for o in r["ops"]:
                    self.assertTrue(all(ist[p] <= ist[o["op"]] for p in o["oncul"]), (c, y, o["op"]))
                self.assertTrue(all(i["sure"] <= c for i in r["istasyonlar"]))
                self.assertGreaterEqual(r["n"], r["teorik"])
        self.assertEqual(main.dengele(ops(), D(274))["n"], 1)
        self.assertEqual((main.dengele(ops(), D(50))["n"], main.dengele(ops(), D(50))["teorik"]), (6, 6))

    def test_lct_ve_cevrim_etkisi(self):
        self.assertEqual(atama(main.dengele(ops(), D(80), "lct")), [["A", "B", "E"], ["C", "F"], ["D", "H", "G"], ["I", "J"]])
        r = main.dengele(ops(), D(100))
        self.assertEqual((r["n"], r["teorik"], r["verim"]), (3, 3, D("91.33")))
        self.assertEqual(atama(r), [["A", "C"], ["B", "D", "F"], ["G", "E", "H", "I", "J"]])        # E, öncülü B atanmadan ilk istasyona giremez

    def test_hatalar(self):
        with self.assertRaisesRegex(ValueError, "uzun operasyon var: C"):
            main.dengele(ops(), D(45))
        with self.assertRaisesRegex(ValueError, "öncülü listede yok: Z"):
            main.dengele(ops([("A", 5, "Z")]), D(10))
        with self.assertRaisesRegex(ValueError, "döngü"):
            main.dengele(ops([("A", 5, "C"), ("B", 5, "A"), ("C", 5, "B")]), D(10))
        with self.assertRaises(ValueError):
            main.dengele([], D(10))


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "h.xlsx"
            r = main.calistir(ORNEK / "operasyonlar.csv", cikti, D(80))
            self.assertEqual((atama(r), r["verim"], r["uyarilar"]), ([["A", "B", "E"], ["C", "F"], ["D", "G", "H"], ["I", "J"]], D("85.63"), []))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "İstasyonlar", "Operasyonlar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "o.csv"
            y.write_text("Operasyon;Süre (sn);Öncül\nA;12,5;-\nB;7.5;A\nA;3;\nC;yok;\nD;4;A; B\n", encoding="utf-8")
            o, uy = main.oku(y)
            self.assertEqual([(x["op"], x["sure"], x["oncul"]) for x in o], [("A", D("12.5"), []), ("B", D("7.5"), ["A"]), ("D", D(4), ["A"])])
            self.assertEqual(sorted(u["tur"] for u in uy), ["Operasyon tekrar ediyor", "Süre okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--operasyonlar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--cevrim", "20", "--yontem", "lct"]), 0)
            self.assertEqual(main.main(arg + ["--talep", "1000", "--sure-dk", "400"]), 0)             # çevrim 24 sn
            self.assertEqual(main.main(arg), 2)
            self.assertEqual(main.main(arg + ["--cevrim", "0"]), 2)
            self.assertEqual(main.main(arg + ["--cevrim", "10"]), 1)                                  # A çevrimden uzun
            self.assertEqual(main.main(["--operasyonlar", str(Path(tmp) / "yok.csv")]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
