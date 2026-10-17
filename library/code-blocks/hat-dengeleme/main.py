"""
Hat Dengeleme — Gerisi Bende kod bloğu
Endüstri Mühendisliği ve Süreç Geliştirme › Endüstri Mühendisi

Operasyon süreleri ve öncelik ilişkilerinden hedef çevrim süresine göre istasyon atamasını ve hat verimini hesaplar.
  Çevrim süresi (C)       = --cevrim;  verilmezse net çalışma süresi (--sure-dk × 60) ÷ talep (--talep)
  Teorik asgari istasyon  = ⌈ Σ operasyon süresi ÷ C ⌉
  Atama (rpw)             = konum ağırlığı (kendi süresi + kendisinden sonra gelen tüm operasyonların süresi) büyükten küçüğe sıralanır;
                            öncülleri atanmış ve kalan süreye sığan ilk operasyon istasyona eklenir, sığan kalmayınca yeni istasyon açılır
  Atama (lct)             = aynı kural, sıralama operasyon süresine göre (en uzun önce)
  Hat verimi %            = Σ süre ÷ (istasyon sayısı × C) × 100;   denge gecikmesi = 100 − hat verimi
  Düzgünlük endeksi       = √ Σ (en dolu istasyon süresi − istasyon süresi)²
Sezgisel bir yöntemdir; en az istasyonlu çözümü garanti etmez. Yerleşim, ekipman, beceri ve bölge kısıtlarını bilmez.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"op": ("operasyon", "operasyon kodu", "is elemani", "kod"), "ad": ("aciklama", "operasyon adi", "tanim"), "sure": ("sure sn", "sure", "standart sure", "sure dk"),
         "oncul": ("oncul", "onculler", "onceki operasyonlar", "oncelik")}
Z = Decimal(0)


def ondalik(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    try:
        return Decimal(re.sub(r"[^\d,.\-]", "", str(x)).replace(",", "."))
    except InvalidOperation:
        return None


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, ops, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("op", "sure")):
        kod = metin(r["op"])
        if not kod:
            continue
        s = ondalik(r["sure"])
        if s is None or s <= 0:
            uy.append(uyari("Yüksek", "Süre okunamadı", kod, f"{yol.name} satır {r['_satir']}: '{metin(r['sure'])}'. Operasyon dengelemeye alınmadı."))
            continue
        if katla(kod) in gorulen:
            uy.append(uyari("Yüksek", "Operasyon tekrar ediyor", kod, f"{yol.name} satır {r['_satir']}: aynı kod iki kez; ilk satır kullanıldı."))
            continue
        gorulen.add(katla(kod))
        ops.append({"op": kod, "ad": metin(r["ad"]), "sure": s, "oncul": [x.strip() for x in re.split(r"[,;|\s]+", metin(r["oncul"])) if x.strip() and x.strip() not in ("-", "—")]})
    return ops, uy


def dengele(ops: list[dict], cevrim: Decimal, yontem: str = "rpw") -> dict:
    if not ops:
        raise ValueError("Dengelenecek operasyon yok.")
    d = {katla(o["op"]): o for o in ops}
    for o in ops:
        eksik = [x for x in o["oncul"] if katla(x) not in d]
        if eksik:
            raise ValueError(f"'{o['op']}' operasyonunun öncülü listede yok: {', '.join(eksik)}")
        o["onc"] = [katla(x) for x in o["oncul"]]
    uzun = [o["op"] for o in ops if o["sure"] > cevrim]
    if uzun:
        raise ValueError(f"Süresi çevrim süresinden ({cevrim:f}) uzun operasyon var: {', '.join(uzun)}. Operasyonu bölün, paralel istasyon kurun veya çevrim süresini artırın.")
    ardil = {k: set() for k in d}
    for k, o in d.items():
        for p in o["onc"]:
            ardil[p].add(k)

    def tum_ardillar(k, yol=()):
        if k in yol:
            raise ValueError(f"Öncelik ilişkilerinde döngü var: {' → '.join(d[x]['op'] for x in (*yol[yol.index(k):], k))}")
        s = set()
        for a in ardil[k]:
            s |= {a} | tum_ardillar(a, (*yol, k))
        return s

    for i, (k, o) in enumerate(d.items()):
        o["agirlik"] = o["sure"] + sum((d[a]["sure"] for a in tum_ardillar(k)), Z)
        o["sira0"] = i
    sirali = sorted(d, key=lambda k: (-(d[k]["agirlik"] if yontem == "rpw" else d[k]["sure"]), -d[k]["sure"], d[k]["sira0"]))
    atanan, istasyonlar = {}, []
    while len(atanan) < len(d):
        ist, kalan = [], cevrim
        while True:
            k = next((k for k in sirali if k not in atanan and d[k]["sure"] <= kalan and all(p in atanan for p in d[k]["onc"])), None)
            if k is None:
                break
            atanan[k] = len(istasyonlar) + 1
            ist.append(k)
            kalan -= d[k]["sure"]
        istasyonlar.append({"no": len(istasyonlar) + 1, "ops": [d[k] for k in ist], "sure": cevrim - kalan, "bos": kalan, "doluluk": yuvarla((cevrim - kalan) / cevrim * 100)})
    for k, o in d.items():
        o["istasyon"] = atanan[k]
    toplam, n = sum((o["sure"] for o in ops), Z), len(istasyonlar)
    tmax = max(i["sure"] for i in istasyonlar)
    verim = yuvarla(toplam / (n * cevrim) * 100)
    return {"ops": ops, "istasyonlar": istasyonlar, "cevrim": cevrim, "toplam": toplam, "n": n, "teorik": int((toplam / cevrim).to_integral_value(rounding="ROUND_CEILING")), "verim": verim,
            "gecikme": 100 - verim, "duzgunluk": yuvarla(sum(((tmax - i["sure"]) ** 2 for i in istasyonlar), Z).sqrt()), "gerceklesen": tmax, "saatlik": yuvarla(Decimal(3600) / cevrim),
            "bos": sum((i["bos"] for i in istasyonlar), Z), "yontem": yontem}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    _baslik(ws, ["Gösterge", "Değer", "Açıklama"], (34, 16, 80))
    for ad, v, ac in (("Çevrim süresi", r["cevrim"], "Her istasyonun bir parça için kullanabileceği en uzun süre"), ("Toplam operasyon süresi", r["toplam"], ""),
                      ("Operasyon sayısı", len(r["ops"]), ""), ("Teorik asgari istasyon", r["teorik"], "⌈ toplam süre ÷ çevrim süresi ⌉"),
                      ("Bulunan istasyon sayısı", r["n"], "Teorik asgariye eşitse daha iyisi yoktur" if r["n"] == r["teorik"] else "Teorik asgarinin üstünde; daha iyi bir çözüm olabilir veya mümkün olmayabilir"),
                      ("Hat verimi %", r["verim"], "Toplam süre ÷ (istasyon sayısı × çevrim süresi)"), ("Denge gecikmesi %", r["gecikme"], "100 − hat verimi"),
                      ("Toplam boş süre", r["bos"], "İstasyonlardaki kullanılmayan süre (çevrim başına)"), ("Düzgünlük endeksi", r["duzgunluk"], "0 = tüm istasyonlar eşit yüklü"),
                      ("En dolu istasyon süresi", r["gerceklesen"], "Hattın bu atamayla çalışabileceği en kısa çevrim"),
                      ("Saatlik çıktı (çevrim süresi saniye ise)", r["saatlik"], "3600 ÷ çevrim süresi"), ("Yöntem", "Konum ağırlığı (RPW)" if r["yontem"] == "rpw" else "En uzun süre önce (LCT)", "")):
        _satir(ws, [ad, v, ac], para_s=(2,) if isinstance(v, Decimal) else ())
    ws = wb.create_sheet("İstasyonlar")
    _baslik(ws, ["İstasyon", "Operasyonlar", "Operasyon Sayısı", "İstasyon Süresi", "Boş Süre", "Doluluk %"], (10, 60, 10, 14, 12, 11))
    for i in r["istasyonlar"]:
        _satir(ws, [i["no"], ", ".join(o["op"] for o in i["ops"]), len(i["ops"]), i["sure"], i["bos"], i["doluluk"]], para_s=(4, 5, 6), dolgu={6: SARI if i["doluluk"] < 70 else None}, sar=(2,))
    ws = wb.create_sheet("Operasyonlar")
    _baslik(ws, ["Operasyon", "Açıklama", "Süre", "Öncüller", "Konum Ağırlığı", "İstasyon"], (14, 40, 10, 20, 13, 10))
    for o in sorted(r["ops"], key=lambda o: (o["istasyon"], -o["agirlik"])):
        _satir(ws, [o["op"], o["ad"], o["sure"], ", ".join(o["oncul"]), o["agirlik"], o["istasyon"]], para_s=(3, 5))
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, cevrim: Decimal, yontem: str = "rpw") -> dict:
    ops, uy = oku(yol)
    r = dengele(ops, cevrim, yontem)
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Operasyonları hedef çevrim süresine göre istasyonlara atar ve hat verimini hesaplar.")
    p.add_argument("--operasyonlar", type=Path, default=ORNEK / "operasyonlar.csv", help="Operasyon, Açıklama, Süre (sn), Öncül (virgülle)")
    p.add_argument("--cevrim", help="Hedef çevrim süresi (operasyon süreleriyle aynı birimde; örnek veride 80)")
    p.add_argument("--talep", help="Vardiyalık talep (adet); --sure-dk ile birlikte çevrim süresini hesaplar")
    p.add_argument("--sure-dk", help="Net vardiya süresi (dk)")
    p.add_argument("--yontem", choices=("rpw", "lct"), default="rpw", help="rpw = konum ağırlığı (varsayılan), lct = en uzun süre önce")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "hat_dengeleme.xlsx")
    a = p.parse_args(argv)
    if not a.operasyonlar.exists():
        print(f"[X] Dosya bulunamadı: {a.operasyonlar}")
        return 1
    c, talep, sure = (para(x) if x else None for x in (a.cevrim, a.talep, a.sure_dk))
    if c is None and talep and sure and talep > 0 and sure > 0:
        c = sure * 60 / talep
    if c is None and not (a.cevrim or a.talep or a.sure_dk) and a.operasyonlar == ORNEK / "operasyonlar.csv":
        c = Decimal(80)
    if c is None or c <= 0:
        print("[X] Çevrim süresi gerekli: --cevrim 80  veya  --talep 330 --sure-dk 440 (süreler saniye ise)")
        return 2
    try:
        r = calistir(a.operasyonlar, a.cikti, c, a.yontem)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] Çevrim {yuvarla(r['cevrim'])} · {len(r['ops'])} operasyon · {r['n']} istasyon (teorik asgari {r['teorik']}) · hat verimi %{r['verim']}")
    for i in r["istasyonlar"]:
        print(f"     İstasyon {i['no']}: {', '.join(o['op'] for o in i['ops'])} · süre {i['sure']:f} · doluluk %{i['doluluk']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
