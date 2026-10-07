import sys
import tempfile
import unittest
from pathlib import Path

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))

import main  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402


class AlgoritmaTesti(unittest.TestCase):
    # Bilinen geçerli örnekler: python-stdnum belgeleri (VKN, TCKN) ve IBAN kayıt örneği (TR)
    def test_vkn(self):
        self.assertTrue(main.vkn_gecerli_mi("4540536920")[0])
        self.assertFalse(main.vkn_gecerli_mi("4540536921")[0])
        self.assertFalse(main.vkn_gecerli_mi("454053692")[0])

    def test_tckn(self):
        self.assertTrue(main.tckn_gecerli_mi("17291716060")[0])
        self.assertFalse(main.tckn_gecerli_mi("17291716050")[0])
        self.assertEqual(main.tckn_gecerli_mi("07291716092"), (False, "İlk hane 0 olamaz"))

    def test_iban(self):
        self.assertTrue(main.iban_gecerli_mi("TR330006100519786457841326")[0])
        self.assertFalse(main.iban_gecerli_mi("TR330006100519786457841327")[0])
        self.assertFalse(main.iban_gecerli_mi("TR33000610051978645784132")[0])   # 25 karakter
        self.assertTrue(main.iban_gecerli_mi("DE89370400440532013000")[0])        # yabancı IBAN

    def test_normallestirme(self):
        s = main.dogrula(" 454 053 6920 ")
        self.assertEqual((s.tur, s.normal, s.gecerli), ("VKN", "4540536920", True))
        s = main.dogrula("tr33 0006 1005 1978 6457 8413 26", "IBAN")
        self.assertTrue(s.gecerli)

    def test_excel_bastaki_sifir(self):
        # 0830166136 Excel'de sayı olarak 830166136 görünür
        s = main.dogrula(830166136, "VKN/TCKN")
        self.assertEqual(s.normal, "0830166136")
        self.assertIn("0 eklendi", s.aciklama)


class UctanUcaTest(unittest.TestCase):
    def test_ornek_csv(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "d.xlsx"
            sonuc = main.calistir(KLASOR / "ornek_veri" / "cari_kartlar.csv", cikti)
            self.assertEqual(sonuc["sutunlar"], ["Vergi No", "IBAN"])
            self.assertEqual(sonuc["satir"], 10)
            wb = load_workbook(cikti)
            hatalar = {(r[1], r[0]) for r in wb["Hatalılar"].iter_rows(min_row=2, values_only=True)}
            # örnek veride bilerek bozulanlar: VKN satır 4, TCKN satır 7, IBAN satır 5, 8 ve 11
            self.assertEqual(hatalar, {("Vergi No", 4), ("Vergi No", 7), ("IBAN", 5), ("IBAN", 8), ("IBAN", 11)})
            ws = wb["Doğrulama"]
            basliklar = [c.value for c in ws[1]]
            sutun = basliklar.index("Vergi No · Açıklama")
            aciklamalar = [r[sutun] or "" for r in ws.iter_rows(min_row=2, values_only=True)]
            self.assertEqual(sum("Mükerrer" in a for a in aciklamalar), 2)   # satır 2 ve 10 aynı VKN

    def test_xlsx_girdi_ve_sutun_secimi(self):
        with tempfile.TemporaryDirectory() as tmp:
            girdi = Path(tmp) / "g.xlsx"
            wb = Workbook()
            wb.active.append(["Ad", "Kimlik"])
            wb.active.append(["A", 17291716060])
            wb.active.append(["B", 830166136])
            wb.save(girdi)
            sonuc = main.calistir(girdi, Path(tmp) / "c.xlsx", ["Kimlik"])
            self.assertEqual(sonuc["hatali"], 0)


if __name__ == "__main__":
    unittest.main()
