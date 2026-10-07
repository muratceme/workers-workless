"""
VKN / TCKN / IBAN Doğrulama — Workers / Workless kod bloğu
Muhasebe › Ön Muhasebe Elemanı (ayrıca: Gişe Yetkilisi)

Cari kart, personel veya ödeme listelerindeki vergi kimlik numaralarını (VKN), T.C. kimlik
numaralarını (TCKN) ve IBAN'ları algoritmik olarak doğrular; hatalı, eksik haneli ve mükerrer
kayıtları işaretleyen bir Excel raporu üretir. İnternete bağlanmaz: numaranın gerçekten birine
ait olup olmadığını değil, matematiksel olarak geçerli olup olmadığını kontrol eder.

Kullanım:
    python main.py                                       # örnek veriyle dener
    python main.py --girdi cari_kartlar.xlsx --cikti cikti/dogrulama.xlsx
    python main.py --girdi liste.csv --sutun "Vergi No" --sutun "IBAN"
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent

# ----------------------------------------------------------------------------
# Algoritmalar
# ----------------------------------------------------------------------------

def tckn_gecerli_mi(no: str) -> tuple[bool, str]:
    """T.C. Kimlik No: 11 hane, ilk hane 0 olamaz.
    10. hane = ((1+3+5+7+9. haneler) × 7 − (2+4+6+8. haneler)) mod 10
    11. hane = (ilk 10 hanenin toplamı) mod 10"""
    if not re.fullmatch(r"\d{11}", no):
        return False, f"11 hane olmalı ({len(no)} hane)"
    if no[0] == "0":
        return False, "İlk hane 0 olamaz"
    d = [int(c) for c in no]
    h10 = (sum(d[0:9:2]) * 7 - sum(d[1:8:2])) % 10
    h11 = sum(d[:10]) % 10
    if d[9] != h10 or d[10] != h11:
        return False, "Kontrol haneleri tutmuyor"
    return True, ""


def vkn_gecerli_mi(no: str) -> tuple[bool, str]:
    """Vergi Kimlik No: 10 hane, son hane kontrol hanesi.
    Sağdan i. hane (i=1..9) için c1 = (hane + i) mod 10; c1 ≠ 0 ise c2 = (c1 × 2^i) mod 9 (0 ise 9);
    kontrol hanesi = (10 − Σc2 mod 10) mod 10"""
    if not re.fullmatch(r"\d{10}", no):
        return False, f"10 hane olmalı ({len(no)} hane)"
    s = 0
    for i, n in enumerate(reversed(no[:9]), 1):
        c1 = (int(n) + i) % 10
        if c1:
            s += (c1 * 2 ** i) % 9 or 9
    if (10 - s % 10) % 10 != int(no[9]):
        return False, "Kontrol hanesi tutmuyor"
    return True, ""


def iban_gecerli_mi(iban: str) -> tuple[bool, str]:
    """ISO 13616 IBAN: ilk 4 karakter sona alınır, harfler 10-35'e çevrilir, sayı mod 97 = 1 olmalı.
    TR IBAN: 26 karakter = TR + 2 kontrol + 5 banka kodu + 1 rezerv (0) + 16 hesap no."""
    if not re.fullmatch(r"[A-Z]{2}\d{2}[A-Z0-9]{10,30}", iban):
        return False, "IBAN biçimi hatalı"
    if iban.startswith("TR"):
        if len(iban) != 26:
            return False, f"TR IBAN 26 karakter olmalı ({len(iban)})"
        if not iban[2:].isdigit():
            return False, "TR IBAN'da harf olamaz"
        if iban[9] != "0":
            return False, "TR IBAN'ın 10. karakteri (rezerv alan) 0 olmalı"
    sayi = "".join(str(int(c, 36)) for c in iban[4:] + iban[:4])
    if int(sayi) % 97 != 1:
        return False, "IBAN kontrol hanesi (mod 97) tutmuyor"
    return True, ""


# ----------------------------------------------------------------------------
# Normalleştirme ve tür tespiti
# ----------------------------------------------------------------------------

def normallestir(deger) -> tuple[str, list[str]]:
    """Hücre değerini karşılaştırılabilir metne çevirir; yapılan düzeltmeleri bildirir."""
    notlar = []
    if deger is None:
        return "", notlar
    if isinstance(deger, float) and deger.is_integer():
        deger = int(deger)
    s = str(deger).strip().upper()
    temiz = re.sub(r"[\s.\-_/]", "", s)
    if temiz != s and temiz:
        notlar.append("Boşluk/ayraç temizlendi")
    if isinstance(deger, int) and len(temiz) == 9:
        temiz = "0" + temiz
        notlar.append("Excel baştaki 0'ı silmiş olabilir; 0 eklendi")
    return temiz, notlar


def tur_tespit(no: str) -> str:
    if re.fullmatch(r"[A-Z]{2}\d{2}[A-Z0-9]+", no):
        return "IBAN"
    if re.fullmatch(r"\d{10}", no):
        return "VKN"
    if re.fullmatch(r"\d{11}", no):
        return "TCKN"
    return "Bilinmiyor"


DOGRULAYICI = {"VKN": vkn_gecerli_mi, "TCKN": tckn_gecerli_mi, "IBAN": iban_gecerli_mi}


@dataclass
class Sonuc:
    tur: str
    normal: str
    gecerli: bool | None
    aciklama: str


def dogrula(deger, beklenen: str | None = None) -> Sonuc:
    no, notlar = normallestir(deger)
    if not no:
        return Sonuc("", "", None, "Boş")
    tur = tur_tespit(no)
    if beklenen == "IBAN" and tur != "IBAN":
        return Sonuc("IBAN", no, False, "IBAN biçiminde değil")
    if beklenen == "VKN/TCKN" and tur not in {"VKN", "TCKN"}:
        return Sonuc("?", no, False, "10 (VKN) veya 11 (TCKN) haneli olmalı")
    if tur == "Bilinmiyor":
        return Sonuc(tur, no, False, "VKN, TCKN veya IBAN biçiminde değil")
    gecerli, neden = DOGRULAYICI[tur](no)
    return Sonuc(tur, no, gecerli, "; ".join([neden] + notlar if neden else notlar))


# ----------------------------------------------------------------------------
# Dosya okuma / yazma
# ----------------------------------------------------------------------------

ANAHTAR = {
    "VKN/TCKN": re.compile(r"vkn|vergi\s*(kimlik)?\s*(no|numara)|tckn|t\.?c\.?\s*(kimlik)?\s*(no|numara)|kimlik\s*no", re.I),
    "IBAN": re.compile(r"iban", re.I),
}


def tablo_oku(yol: Path) -> tuple[list[str], list[list]]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()  # salt okunur modda dosya kapatılmazsa Windows'ta kilitli kalır
    else:
        metin = None
        for kodlama in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kodlama)
                break
            except UnicodeDecodeError:
                continue
        ayirici = csv.Sniffer().sniff(metin.splitlines()[0], delimiters=",;\t").delimiter
        # CSV'de sayılar metin olarak kalır; baştaki sıfırlar korunur
        satirlar = list(csv.reader(metin.splitlines(), delimiter=ayirici))
    satirlar = [r for r in satirlar if any(c not in (None, "") for c in r)]
    if not satirlar:
        raise SystemExit(f"{yol} boş.")
    genislik = max(len(r) for r in satirlar)
    satirlar = [list(r) + [None] * (genislik - len(r)) for r in satirlar]
    return [str(b or "").strip() for b in satirlar[0]], satirlar[1:]


def sutunlari_bul(basliklar: list[str], istenen: list[str] | None) -> dict[int, str]:
    if istenen:
        bulunan = {}
        for ad in istenen:
            if ad not in basliklar:
                raise SystemExit(f"'{ad}' sütunu bulunamadı. Başlıklar: {basliklar}")
            i = basliklar.index(ad)
            bulunan[i] = "IBAN" if ANAHTAR["IBAN"].search(ad) else None
        return bulunan
    bulunan = {}
    for i, b in enumerate(basliklar):
        for tur, desen in ANAHTAR.items():
            if desen.search(b):
                bulunan[i] = tur
                break
    if not bulunan:
        raise SystemExit("VKN/TCKN veya IBAN sütunu otomatik bulunamadı. --sutun ile belirtin. "
                         f"Başlıklar: {basliklar}")
    return bulunan


KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")


def calistir(girdi: Path, cikti: Path, sutunlar: list[str] | None = None) -> dict:
    basliklar, satirlar = tablo_oku(girdi)
    hedef = sutunlari_bul(basliklar, sutunlar)

    sonuclar = {i: [dogrula(r[i], hedef[i]) for r in satirlar] for i in hedef}
    tekrar = {i: Counter(s.normal for s in sonuclar[i] if s.normal) for i in hedef}

    wb = Workbook()
    ws = wb.active
    ws.title = "Doğrulama"
    yeni_basliklar = list(basliklar)
    for i in hedef:
        yeni_basliklar += [f"{basliklar[i]} · Tür", f"{basliklar[i]} · Normal", f"{basliklar[i]} · Durum", f"{basliklar[i]} · Açıklama"]
    ws.append(yeni_basliklar)
    for h in ws[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI

    hatalar = []
    ozet = Counter()
    for n, r in enumerate(satirlar):
        ek = []
        satir_hatali = satir_uyari = False
        for i in hedef:
            s = sonuclar[i][n]
            aciklama = s.aciklama
            if s.normal and tekrar[i][s.normal] > 1:
                aciklama = "; ".join(filter(None, [aciklama, f"Mükerrer ({tekrar[i][s.normal]} kez)"]))
                satir_uyari = True
            durum = "Boş" if s.gecerli is None else ("Geçerli" if s.gecerli else "HATALI")
            ozet[(basliklar[i], s.tur or "-", durum)] += 1
            if s.gecerli is False:
                satir_hatali = True
                hatalar.append([n + 2, basliklar[i], r[i], s.tur, s.aciklama])
            ek += [s.tur, s.normal, durum, aciklama]
        ws.append([*r, *ek])
        if satir_hatali or satir_uyari:
            for h in ws[ws.max_row]:
                h.fill = KIRMIZI if satir_hatali else SARI
    for i in range(1, len(yeni_basliklar) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 18
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    hs = wb.create_sheet("Hatalılar")
    hs.append(["Excel Satırı", "Sütun", "Değer", "Tür", "Neden"])
    for h in hs[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for x in hatalar:
        hs.append(x)
    for c, g in zip("ABCDE", (12, 20, 30, 10, 50)):
        hs.column_dimensions[c].width = g

    oz = wb.create_sheet("Özet")
    oz.append(["Sütun", "Tür", "Durum", "Adet"])
    for h in oz[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for (sutun, tur, durum), adet in sorted(ozet.items()):
        oz.append([sutun, tur, durum, adet])
    for c, g in zip("ABCD", (24, 12, 12, 10)):
        oz.column_dimensions[c].width = g

    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    return {"satir": len(satirlar), "hatali": len(hatalar), "sutunlar": [basliklar[i] for i in hedef]}


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="VKN, TCKN ve IBAN'ları toplu doğrular.")
    p.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "cari_kartlar.csv", help="Liste (.xlsx/.csv)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "dogrulama.xlsx")
    p.add_argument("--sutun", action="append", help="Kontrol edilecek sütun başlığı (birden çok verilebilir)")
    a = p.parse_args(argv)
    sonuc = calistir(a.girdi, a.cikti, a.sutun)
    print(f"[OK] {sonuc['satir']} satır, {', '.join(sonuc['sutunlar'])} sütunları kontrol edildi")
    print(f"[{'!' if sonuc['hatali'] else 'OK'}] {sonuc['hatali']} hatalı değer")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
