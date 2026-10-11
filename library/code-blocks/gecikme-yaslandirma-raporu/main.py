"""
Gecikme Yaşlandırma Raporu — Workers / Workless kod bloğu
Krediler İzleme ve Takip › Kredi İzleme Uzmanı

Kredileri gecikme gün aralıklarına (Güncel, 1-30, 31-60, 61-90, 90+) dağıtır; şube ve ürün bazında raporlar.
  Gecikme günü  = rapor tarihi − ödenmemiş en eski taksitin vadesi   ("Gecikme Gün" sütunu doluysa o kullanılır)
  PAR30 / PAR90 = 30 (90) günden fazla gecikmiş kredilerin TOPLAM risk bakiyesi ÷ portföyün toplam risk bakiyesi
                  (yalnız geciken taksit değil, kredinin tüm bakiyesi sayılır)
  Müşteri bazı  = müşterinin en çok gecikmiş kredisinin aralığı; müşterinin tüm bakiyesi o aralıkta gösterilir.
  Geçiş matrisi = önceki dönem dosyası verilirse: bir önceki aralıktan bu döneme hangi aralığa geçildi (bakiye bazında);
                  kötüleşen, iyileşen, yeni giren ve kapanan krediler.
Aralık sınırları --araliklar ile değiştirilebilir (varsayılan 30,60,90).
Not: Rapor bir sınıflandırma kararı değildir. BDDK Karşılıklar Yönetmeliği'nde 30 günden fazla gecikme yakın izleme,
90 günden fazla gecikme donuk alacak göstergesidir; sınıflandırma başka ölçütleri de içerir. İnternete bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: 60 kredi, önceki ayla karşılaştırmalı
    python main.py --krediler krediler.xlsx --onceki gecen_ay.xlsx --tarih 30.09.2026 --araliklar 30,60,90,180
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


KREDI_SUTUNLARI = {"no": ("kredi no", "kredi numarasi", "hesap no"), "musteri_no": ("musteri no", "musteri numarasi"), "musteri": ("musteri", "musteri adi", "unvan", "ad soyad"),
                   "sube": ("sube",), "urun": ("urun", "kredi turu", "urun turu"), "bakiye": ("risk bakiyesi", "bakiye", "anapara bakiyesi", "kalan anapara"),
                   "geciken": ("geciken tutar", "gecikmis tutar", "odenmeyen tutar"), "vade": ("en eski vade", "en eski vade tarihi", "en eski odenmemis vade", "ilk gecikme tarihi"),
                   "gun": ("gecikme gun", "gecikme gunu", "gecikme gun sayisi", "gecikme")}


@dataclass
class Kredi:
    no: str
    musteri_no: str
    musteri: str
    sube: str
    urun: str
    bakiye: Decimal
    geciken: Decimal
    gun: int
    aralik: str = ""


def aralik_adlari(sinirlar: list[int]) -> list[str]:
    return ["Güncel"] + [f"{a + 1}-{b}" for a, b in zip([0] + sinirlar, sinirlar)] + [f"{sinirlar[-1]}+"]


def aralik_bul(gun: int, sinirlar: list[int]) -> str:
    adlar = aralik_adlari(sinirlar)
    if gun <= 0:
        return adlar[0]
    return next((adlar[i + 1] for i, s in enumerate(sinirlar) if gun <= s), adlar[-1])


def oku(yol: Path, rapor: date, sinirlar: list[int]) -> tuple[list[Kredi], list[dict]]:
    uy, krediler, gorulen = [], [], set()
    for r in kayitlar(yol, KREDI_SUTUNLARI, ("no", "bakiye")):
        no = metin(r["no"])
        if not no:
            continue
        if no in gorulen:
            uy.append(uyari("Yüksek", "Mükerrer kredi", no, f"{yol.name} satır {r['_satir']}: kredi numarası tekrar ediyor; ilk satır kullanıldı."))
            continue
        b, g = para(r["bakiye"]), para(r["geciken"]) or Decimal(0)
        if b is None or b < 0:
            uy.append(uyari("Yüksek", "Bakiye okunamadı", no, f"{yol.name} satır {r['_satir']}: '{metin(r['bakiye'])}'. Kredi rapora alınmadı."))
            continue
        gun, v = para(r["gun"]), tarih(r["vade"])
        if gun is None:
            if metin(r["vade"]) and not v:
                uy.append(uyari("Yüksek", "Vade tarihi okunamadı", no, f"{yol.name} satır {r['_satir']}: '{metin(r['vade'])}'. Kredi güncel sayıldı; kontrol edin."))
            gun = Decimal((rapor - v).days) if v else Decimal(0)
        gun = max(int(gun), 0)
        if g > 0 and gun == 0:
            uy.append(uyari("Orta", "Geciken tutar var, gecikme günü yok", no, f"Geciken tutar {g:f}, gecikme günü 0 (vade tarihi boş veya ileri tarihli). Kredi güncel sayıldı."))
        if gun > 0 and g == 0:
            uy.append(uyari("Orta", "Gecikme günü var, geciken tutar yok", no, f"{gun} gün gecikme görünüyor ama geciken tutar 0."))
        if g > b:
            uy.append(uyari("Orta", "Geciken tutar bakiyeden büyük", no, f"Geciken {g:f} > risk bakiyesi {b:f}. Faiz ve masraf dahil olabilir; tanımı kontrol edin."))
        gorulen.add(no)
        krediler.append(Kredi(no, metin(r["musteri_no"]) or metin(r["musteri"]) or no, metin(r["musteri"]), metin(r["sube"]) or "Belirtilmemiş", metin(r["urun"]) or "Belirtilmemiş",
                              b, g, gun, aralik_bul(gun, sinirlar)))
    return krediler, uy


def dagilim(krediler: list[Kredi], adlar: list[str], anahtar=None) -> list[dict]:
    """anahtar verilirse (ör. şube) grup × aralık; verilmezse yalnız aralık."""
    toplam = sum((k.bakiye for k in krediler), Decimal(0))
    gruplar = sorted({anahtar(k) for k in krediler}) if anahtar else [None]
    sonuc = []
    for g in gruplar:
        l = [k for k in krediler if not anahtar or anahtar(k) == g]
        gt = sum((k.bakiye for k in l), Decimal(0))
        s = {"grup": g, "adet": len(l), "bakiye": gt, "araliklar": {}}
        for a in adlar:
            la = [k for k in l if k.aralik == a]
            ba = sum((k.bakiye for k in la), Decimal(0))
            s["araliklar"][a] = {"adet": len(la), "musteri": len({k.musteri_no for k in la}), "bakiye": ba, "geciken": sum((k.geciken for k in la), Decimal(0)),
                                 "pay": yuvarla(ba / gt * 100) if gt else Decimal(0)}
        s["pay_portfoy"] = yuvarla(gt / toplam * 100) if toplam else Decimal(0)
        sonuc.append(s)
    return sonuc


def par(krediler: list[Kredi], esik: int) -> Decimal:
    toplam = sum((k.bakiye for k in krediler), Decimal(0))
    return yuvarla(sum((k.bakiye for k in krediler if k.gun > esik), Decimal(0)) / toplam * 100) if toplam else Decimal(0)


def analiz(krediler: list[Kredi], sinirlar: list[int], onceki: list[Kredi] | None = None) -> dict:
    adlar = aralik_adlari(sinirlar)
    sira = {a: i for i, a in enumerate(adlar)}
    r = {"adlar": adlar, "krediler": krediler, "genel": dagilim(krediler, adlar)[0], "sube": dagilim(krediler, adlar, lambda k: k.sube),
         "urun": dagilim(krediler, adlar, lambda k: k.urun), "par": {e: par(krediler, e) for e in sorted({30, 90} | {s for s in sinirlar if s in (30, 60, 90)})},
         "gecikmeli_oran": par(krediler, 0), "uyarilar": [], "gecis": None, "hareket": []}
    for veri, f in ((r["sube"], lambda k: k.sube), (r["urun"], lambda k: k.urun)):
        for d in veri:
            l = [k for k in krediler if f(k) == d["grup"]]
            d["par30"], d["par90"] = par(l, 30), par(l, 90)
    # müşteri bazı: en kötü aralık
    musteri = defaultdict(list)
    for k in krediler:
        musteri[k.musteri_no].append(k)
    mb = {a: {"musteri": 0, "bakiye": Decimal(0)} for a in adlar}
    r["musteriler"] = []
    for no, l in musteri.items():
        kotu = max(l, key=lambda k: k.gun)
        mb[kotu.aralik]["musteri"] += 1
        mb[kotu.aralik]["bakiye"] += sum(k.bakiye for k in l)
        if kotu.gun > 0:
            r["musteriler"].append({"no": no, "ad": kotu.musteri, "kredi": len(l), "gecikmeli": sum(k.gun > 0 for k in l), "bakiye": sum(k.bakiye for k in l),
                                    "geciken": sum(k.geciken for k in l), "gun": kotu.gun, "aralik": kotu.aralik,
                                    "bulasma": sum(k.bakiye for k in l if sira[k.aralik] < sira[kotu.aralik])})
    r["musteri_bazi"] = mb
    r["musteriler"].sort(key=lambda m: (-m["gun"], -m["bakiye"]))
    if onceki is not None:
        eski = {k.no: k for k in onceki}
        yeni = {k.no: k for k in krediler}
        m = {a: {b: Decimal(0) for b in adlar + ["Kapandı"]} for a in adlar + ["Yeni"]}
        for no, e in eski.items():
            y = yeni.get(no)
            m[e.aralik][y.aralik if y else "Kapandı"] += e.bakiye              # geçiş, dönem başı bakiyesiyle ölçülür
            if y and sira[y.aralik] != sira[e.aralik]:
                r["hareket"].append({"k": y, "onceki": e.aralik, "yon": "Kötüleşti" if sira[y.aralik] > sira[e.aralik] else "İyileşti", "onceki_gun": e.gun})
            elif not y and e.gun > 0:
                r["hareket"].append({"k": e, "onceki": e.aralik, "yon": "Kapandı / listede yok", "onceki_gun": e.gun})
        for no, y in yeni.items():
            if no not in eski:
                m["Yeni"][y.aralik] += y.bakiye
                if y.gun > 0:
                    r["hareket"].append({"k": y, "onceki": "Yeni", "yon": "Yeni kredi gecikmede", "onceki_gun": None})
        r["gecis"] = m
        r["hareket"].sort(key=lambda h: ({"Kötüleşti": 0, "Yeni kredi gecikmede": 1, "İyileşti": 2}.get(h["yon"], 3), -h["k"].bakiye))
        r["onceki_par"] = {e: par(onceki, e) for e in r["par"]}
    return r


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

def _aralik_renk(i: int, n: int) -> str | None:
    return None if i == 0 else SARI if i < n - 1 else KIRMIZI


def rapor_yaz(cikti: Path, r: dict, rapor: date) -> None:
    wb, adlar = Workbook(), r["adlar"]
    oz = wb.active
    oz.title = "Yaşlandırma"
    oz.append([f"Rapor tarihi {rapor:%d.%m.%Y}"])
    oz["A1"].font = Font(bold=True)
    _baslik(oz, ["Gecikme Aralığı (gün)", "Kredi Adedi", "Müşteri Adedi", "Risk Bakiyesi", "Pay %", "Geciken Tutar", "Müşteri Bazında Müşteri", "Müşteri Bazında Bakiye"],
            (22, 10, 10, 18, 9, 18, 14, 18), "A3")
    g = r["genel"]
    for i, a in enumerate(adlar):
        x, m = g["araliklar"][a], r["musteri_bazi"][a]
        _satir(oz, [a, x["adet"], x["musteri"], x["bakiye"], x["pay"], x["geciken"], m["musteri"], m["bakiye"]], para_s=(4, 5, 6, 8), dolgu={1: _aralik_renk(i, len(adlar))})
    _satir(oz, ["Toplam", g["adet"], len({k.musteri_no for k in r["krediler"]}), g["bakiye"], 100, sum(k.geciken for k in r["krediler"]), "", g["bakiye"]], para_s=(4, 5, 6, 8))
    for c in oz[oz.max_row]:
        c.font = Font(bold=True)
    oz.append([])
    _satir(oz, ["Gecikmeli kredi oranı % (1+ gün)", "", "", "", r["gecikmeli_oran"]], para_s=(5,))
    for e, v in r["par"].items():
        n = _satir(oz, [f"PAR{e} % ({e} günden fazla)", "", "", "", v], para_s=(5,))
        if r.get("onceki_par"):
            oz.cell(n, 6).value = f"önceki dönem %{r['onceki_par'][e]}"
    oz.append([])
    oz.append(["PAR: gecikmiş kredinin yalnız geciken taksiti değil tüm risk bakiyesi sayılır. Müşteri bazı: müşterinin tüm bakiyesi en çok gecikmiş kredisinin aralığında gösterilir."])

    for ad, veri, baslik in (("Şube", r["sube"], "Şube"), ("Ürün", r["urun"], "Ürün")):
        ws = wb.create_sheet(f"{ad} Bazında")
        _baslik(ws, [baslik, "Kredi Adedi", "Risk Bakiyesi", "Portföy Payı %", *adlar, "PAR30 %", "PAR90 %"], (24, 10, 18, 11, *[16] * len(adlar), 10, 10))
        for d in sorted(veri, key=lambda d: -d["par30"]):
            _satir(ws, [d["grup"], d["adet"], d["bakiye"], d["pay_portfoy"], *[d["araliklar"][a]["bakiye"] for a in adlar], d["par30"], d["par90"]],
                   para_s=(3, 4, *range(5, 7 + len(adlar))), dolgu={5 + len(adlar): KIRMIZI if d["par30"] > r["par"][30] * Decimal("1.5") and d["par30"] > 0 else None})
        ws.append([])
        ws.append([f"Kırmızı: PAR30 oranı portföy ortalamasının (%{r['par'][30]}) 1,5 katından yüksek."])

    kl = wb.create_sheet("Gecikmeli Krediler")
    _baslik(kl, ["Kredi No", "Müşteri No", "Müşteri", "Şube", "Ürün", "Risk Bakiyesi", "Geciken Tutar", "Gecikme Gün", "Aralık", "Aksiyon / Not"], (14, 12, 30, 16, 20, 16, 16, 9, 10, 40))
    for k in sorted((k for k in r["krediler"] if k.gun > 0), key=lambda k: (-k.gun, -k.bakiye)):
        _satir(kl, [k.no, k.musteri_no, k.musteri, k.sube, k.urun, k.bakiye, k.geciken, k.gun, k.aralik, ""], para_s=(6, 7),
               dolgu={9: _aralik_renk(adlar.index(k.aralik), len(adlar)), 10: SARI})
    kl.auto_filter.ref = f"A1:J{kl.max_row}"

    mu = wb.create_sheet("Müşteri Bazında")
    _baslik(mu, ["Müşteri No", "Müşteri", "Kredi Adedi", "Gecikmeli Kredi", "Toplam Bakiye", "Geciken Tutar", "En Uzun Gecikme (gün)", "Aralık", "Daha İyi Aralıktaki Bakiye"],
            (12, 30, 9, 10, 16, 16, 12, 10, 18))
    for m in r["musteriler"]:
        _satir(mu, [m["no"], m["ad"], m["kredi"], m["gecikmeli"], m["bakiye"], m["geciken"], m["gun"], m["aralik"], m["bulasma"]], para_s=(5, 6, 9),
               dolgu={8: _aralik_renk(adlar.index(m["aralik"]), len(adlar))})

    if r["gecis"] is not None:
        gm = wb.create_sheet("Geçiş Matrisi")
        sut = adlar + ["Kapandı"]
        _baslik(gm, ["Önceki Dönem ↓ / Bu Dönem →", *sut, "Toplam", "Kötüleşen %"], (26, *[16] * len(sut), 16, 11))
        for i, a in enumerate(adlar + ["Yeni"]):
            satir = r["gecis"][a]
            t = sum(satir.values())
            kotu = sum(satir[b] for j, b in enumerate(adlar) if a != "Yeni" and j > i)
            _satir(gm, [a, *[satir[b] for b in sut], t, yuvarla(kotu / t * 100) if t and a != "Yeni" else None], para_s=tuple(range(2, 4 + len(sut))))
        gm.append([])
        gm.append(["Tutarlar önceki dönemin risk bakiyesidir (Yeni satırı bu dönemin bakiyesi). Köşegenin sağı kötüleşme, solu iyileşme."])
        hr = wb.create_sheet("Aralık Değiştirenler")
        _baslik(hr, ["Yön", "Kredi No", "Müşteri", "Şube", "Ürün", "Risk Bakiyesi", "Önceki Aralık", "Bu Dönem", "Önceki Gün", "Gün"], (22, 14, 30, 16, 20, 16, 12, 12, 9, 8))
        for h in r["hareket"]:
            k = h["k"]
            _satir(hr, [h["yon"], k.no, k.musteri, k.sube, k.urun, k.bakiye, h["onceki"], k.aralik if h["yon"] != "Kapandı / listede yok" else "—", h["onceki_gun"],
                        k.gun if h["yon"] != "Kapandı / listede yok" else None], para_s=(6,),
                   dolgu={1: {"Kötüleşti": KIRMIZI, "Yeni kredi gecikmede": SARI, "İyileşti": YESIL}.get(h["yon"], MAVI)})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, rapor: date, sinirlar: list[int], onceki_yol: Path | None = None, onceki_tarih: date | None = None) -> dict:
    krediler, uy = oku(yol, rapor, sinirlar)
    if not krediler:
        raise ValueError(f"{yol.name}: okunabilir kredi yok")
    onceki = None
    if onceki_yol:
        onceki, uy2 = oku(onceki_yol, onceki_tarih or rapor, sinirlar)
        uy += [u for u in uy2 if u["onem"] == "Yüksek"]
    r = analiz(krediler, sinirlar, onceki)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, rapor)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Kredileri gecikme gün aralıklarına dağıtıp şube ve ürün bazında raporlar.")
    p.add_argument("--krediler", type=Path, default=ORNEK / "krediler.csv",
                   help="Kredi No, Müşteri No, Müşteri, Şube, Ürün, Risk Bakiyesi, Geciken Tutar, En Eski Vade Tarihi veya Gecikme Gün")
    p.add_argument("--onceki", type=Path, help="Önceki dönemin aynı biçimli dosyası (geçiş matrisi için)")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 30.09.2026)")
    p.add_argument("--onceki-tarih", help="Önceki dönem rapor tarihi (önceki dosyada gecikme günü yoksa gerekir)")
    p.add_argument("--araliklar", default="30,60,90", help="Aralık üst sınırları, virgülle (varsayılan 30,60,90)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "gecikme_yaslandirma.xlsx")
    a = p.parse_args(argv)
    ornek = a.krediler == ORNEK / "krediler.csv"
    onceki = a.onceki or (ORNEK / "onceki_donem.csv" if ornek else None)
    for y in (a.krediler, onceki):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    rapor = tarih(a.tarih) if a.tarih else (date(2026, 9, 30) if ornek else date.today())
    ot = tarih(a.onceki_tarih) if a.onceki_tarih else (date(2026, 8, 31) if ornek else None)
    try:
        sinirlar = [int(x) for x in a.araliklar.split(",")]
        if not rapor or sinirlar != sorted(set(sinirlar)) or sinirlar[0] <= 0 or (a.onceki_tarih and not ot):
            raise ValueError
    except ValueError:
        print("[X] --tarih GG.AA.YYYY biçiminde; --araliklar artan pozitif sayılar olmalı (ör. 30,60,90)")
        return 2
    try:
        r = calistir(a.krediler, a.cikti, rapor, sinirlar, onceki, ot)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    g = r["genel"]
    print(f"[OK] {g['adet']} kredi · gecikmeli {sum(k.gun > 0 for k in r['krediler'])} · " + " · ".join(f"PAR{e} %{v}" for e, v in r["par"].items()))
    if r["gecis"] is not None:
        print(f"[OK] Önceki döneme göre: {sum(h['yon'] == 'Kötüleşti' for h in r['hareket'])} kötüleşen, {sum(h['yon'] == 'İyileşti' for h in r['hareket'])} iyileşen")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['kim']} · {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
