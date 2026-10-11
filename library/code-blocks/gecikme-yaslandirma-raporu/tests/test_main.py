import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
S = [30, 60, 90]


def kredi(no, bakiye, gun, musteri=None, sube="Merkez", urun="İhtiyaç", geciken=None):
    return main.Kredi(no, musteri or no, "M", sube, urun, D(str(bakiye)), D(str(geciken if geciken is not None else (100 if gun else 0))), gun, main.aralik_bul(gun, S))


class AralikTesti(unittest.TestCase):
    def test_sinirlar(self):
        self.assertEqual(main.aralik_adlari(S), ["Güncel", "1-30", "31-60", "61-90", "90+"])
        self.assertEqual([main.aralik_bul(g, S) for g in (-3, 0, 1, 30, 31, 60, 61, 90, 91, 400)],
                         ["Güncel", "Güncel", "1-30", "1-30", "31-60", "31-60", "61-90", "61-90", "90+", "90+"])
        self.assertEqual(main.aralik_adlari([30, 90, 180]), ["Güncel", "1-30", "31-90", "91-180", "180+"])

    def test_par_ve_dagilim(self):
        k = [kredi("A", 500, 0), kredi("B", 200, 30), kredi("C", 100, 31), kredi("D", 150, 90), kredi("E", 50, 91)]
        r = main.analiz(k, S)
        self.assertEqual((r["par"][30], r["par"][90], r["gecikmeli_oran"]), (D("30.00"), D("5.00"), D("50.00")))   # 300/1000, 50/1000, 500/1000
        a = r["genel"]["araliklar"]
        self.assertEqual([(a[x]["adet"], a[x]["bakiye"], a[x]["pay"]) for x in r["adlar"]],
                         [(1, D(500), D("50.00")), (1, D(200), D("20.00")), (1, D(100), D("10.00")), (1, D(150), D("15.00")), (1, D(50), D("5.00"))])

    def test_grup_bazinda(self):
        k = [kredi("A", 600, 0, sube="X"), kredi("B", 400, 45, sube="X"), kredi("C", 1000, 0, sube="Y", urun="Konut")]
        r = main.analiz(k, S)
        x = next(d for d in r["sube"] if d["grup"] == "X")
        self.assertEqual((x["par30"], x["pay_portfoy"], x["araliklar"]["31-60"]["bakiye"]), (D("40.00"), D("50.00"), D(400)))
        self.assertEqual({d["grup"]: d["par30"] for d in r["urun"]}, {"Konut": D("0.00"), "İhtiyaç": D("40.00")})

    def test_musteri_bazi(self):
        k = [kredi("A", 1000, 0, "M1"), kredi("B", 300, 95, "M1"), kredi("C", 500, 10, "M2"), kredi("D", 200, 0, "M3")]
        r = main.analiz(k, S)
        mb = r["musteri_bazi"]
        self.assertEqual((mb["90+"], mb["1-30"], mb["Güncel"]), ({"musteri": 1, "bakiye": D(1300)}, {"musteri": 1, "bakiye": D(500)}, {"musteri": 1, "bakiye": D(200)}))
        m1 = r["musteriler"][0]
        self.assertEqual((m1["no"], m1["gecikmeli"], m1["bulasma"]), ("M1", 1, D(1000)))        # güncel görünen 1.000 müşteri bazında 90+


class GecisTesti(unittest.TestCase):
    def test_gecis_matrisi(self):
        eski = [kredi("A", 100, 0), kredi("B", 200, 20), kredi("C", 300, 40), kredi("D", 400, 70), kredi("E", 50, 10)]
        yeni = [kredi("A", 90, 15), kredi("B", 190, 50), kredi("C", 280, 0), kredi("D", 390, 100), kredi("F", 500, 5)]
        r = main.analiz(yeni, S, eski)
        g = r["gecis"]
        self.assertEqual((g["Güncel"]["1-30"], g["1-30"]["31-60"], g["31-60"]["Güncel"], g["61-90"]["90+"], g["1-30"]["Kapandı"], g["Yeni"]["1-30"]),
                         (D(100), D(200), D(300), D(400), D(50), D(500)))
        self.assertEqual([(h["yon"], h["k"].no) for h in r["hareket"]],
                         [("Kötüleşti", "D"), ("Kötüleşti", "B"), ("Kötüleşti", "A"), ("Yeni kredi gecikmede", "F"), ("İyileşti", "C"), ("Kapandı / listede yok", "E")])
        self.assertEqual(r["onceki_par"][30], D("66.67"))                                          # (300 + 400) / 1050


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "y.xlsx"
            r = main.calistir(ORNEK / "krediler.csv", cikti, date(2026, 9, 30), S, ORNEK / "onceki_donem.csv", date(2026, 8, 31))
            self.assertEqual(r["genel"]["adet"], 58)
            self.assertEqual(sum(x["bakiye"] for x in r["genel"]["araliklar"].values()), r["genel"]["bakiye"])
            self.assertEqual(sum(sum(s.values()) for a, s in r["gecis"].items() if a != "Yeni"), sum(k.bakiye for k in main.oku(ORNEK / "onceki_donem.csv", date(2026, 8, 31), S)[0]))
            self.assertEqual(sum(m["bakiye"] for m in r["musteri_bazi"].values()), r["genel"]["bakiye"])
            self.assertEqual(load_workbook(cikti).sheetnames,
                             ["Yaşlandırma", "Şube Bazında", "Ürün Bazında", "Gecikmeli Krediler", "Müşteri Bazında", "Geçiş Matrisi", "Aralık Değiştirenler", "Uyarılar"])

    def test_okuma(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "k.csv"
            y.write_text("Kredi No;Risk Bakiyesi;Geciken Tutar;En Eski Vade Tarihi;Gecikme Gün\n"
                         "K1;10.000,00;500,00;21.08.2026;\nK2;5.000,00;200,00;;45\nK3;1.000,00;300,00;;\nK4;2.000,00;0;01.09.2026;\nK5;800,00;900,00;bozuk;\n"
                         "K1;1,00;0;;\nK6;eksi;0;;\n", encoding="utf-8")
            k, uy = main.oku(y, date(2026, 9, 30), S)
            self.assertEqual([(x.no, x.gun, x.aralik) for x in k], [("K1", 40, "31-60"), ("K2", 45, "31-60"), ("K3", 0, "Güncel"), ("K4", 29, "1-30"), ("K5", 0, "Güncel")])
            self.assertEqual(sorted(u["tur"] for u in uy), ["Bakiye okunamadı", "Geciken tutar bakiyeden büyük", "Geciken tutar var, gecikme günü yok",
                                                            "Geciken tutar var, gecikme günü yok", "Gecikme günü var, geciken tutar yok", "Mükerrer kredi", "Vade tarihi okunamadı"])
            arg = ["--krediler", str(y), "--cikti", str(Path(tmp) / "x.xlsx"), "--tarih", "30.09.2026"]
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(load_workbook(Path(tmp) / "x.xlsx").sheetnames[-2:], ["Müşteri Bazında", "Uyarılar"])     # önceki dönem yoksa geçiş sayfası yok
            self.assertEqual(main.main(arg + ["--araliklar", "30,180"]), 0)
            self.assertEqual(main.main(arg + ["--araliklar", "60,30"]), 2)
            self.assertEqual(main.main(["--krediler", str(Path(tmp) / "yok.csv")]), 1)
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "o.xlsx")]), 0)


if __name__ == "__main__":
    unittest.main()
