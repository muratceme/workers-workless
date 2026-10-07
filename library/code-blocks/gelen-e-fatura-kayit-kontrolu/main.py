"""
Gelen e-Fatura Kayıt Kontrolü — Workers / Workless kod bloğu
Muhasebe › Muhasebe Elemanı (ayrıca: Mali Müşavirlik Bürosu › Muhasebe Elemanı)

GİB veya özel entegratör portalından alınan gelen e-Fatura / e-Arşiv listesini (XML/ZIP klasörü ya da
Excel/CSV liste) muhasebe alış kayıtlarıyla karşılaştırır; deftere işlenmemiş, mükerrer işlenmiş,
tutarı farklı ve portalda karşılığı bulunmayan kayıtları tek Excel raporunda verir. İnternete bağlanmaz.

Eşleştirme sırası:
  1. Fatura numarası (16 karakterlik GİB numarası, büyük/küçük harf ve boşluk duyarsız) + satıcı VKN
  2. Fatura numarası tek başına (kayıtta VKN yoksa)
  3. Satıcı VKN + vergiler dahil tutar + tarih farkı ≤ tolerans (fatura numarası yanlış girilmiş kayıtlar)

Kullanım:
    python main.py                                                     # örnek veriyle dener
    python main.py --efatura "C:/Gelen/2026-09" --kayitlar muavin_alislar.xlsx
    python main.py --efatura portal_listesi.xlsx --kayitlar alis_kayitlari.csv --gun-toleransi 5
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import efatura_cekirdek as efatura

BURASI = Path(__file__).resolve().parent


@dataclass
class Belge:
    kaynak: str          # "efatura" | "kayit"
    sira: int
    no: str
    tarih: date | None
    vkn: str
    unvan: str
    tutar: Decimal       # vergiler dahil (ödenecek değil)
    kdv: Decimal | None = None
    aciklama: str = ""

    @property
    def anahtar_no(self) -> str:
        return re.sub(r"[\s\-_/]", "", self.no.upper())


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def para(x) -> Decimal:
    if x in (None, ""):
        return Decimal(0)
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("₺", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return Decimal(0)


def tarih(x) -> date | None:
    if x in (None, ""):
        return None
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    for f in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%d.%m.%Y %H:%M:%S"):
        try:
            return datetime.strptime(str(x).strip(), f).date()
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ayirici = csv.Sniffer().sniff(metin.splitlines()[0], delimiters=",;\t").delimiter
        satirlar = list(csv.reader(metin.splitlines(), delimiter=ayirici))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


ALANLAR = {
    "no": ("fatura no", "fatura numarası", "belge no", "evrak no", "fatura no.", "fiş belge no", "numara"),
    "tarih": ("fatura tarihi", "tarih", "belge tarihi", "düzenleme tarihi"),
    "vkn": ("vkn", "vkn/tckn", "vergi no", "gönderici vkn/tckn", "satıcı vkn", "gönderici vkn", "cari vkn", "vergi numarası"),
    "unvan": ("unvan", "gönderici unvan", "satıcı", "cari unvan", "cari hesap", "firma", "gönderici"),
    "tutar": ("vergiler dahil tutar", "toplam tutar", "genel toplam", "kdv dahil tutar", "tutar", "borç", "alacak"),
    "kdv": ("kdv", "kdv tutarı", "hesaplanan kdv", "toplam kdv"),
    "aciklama": ("açıklama", "fiş açıklaması"),
}


def liste_belgeleri(yol: Path, kaynak: str) -> list[Belge]:
    satirlar = tablo_oku(yol)
    b = [kucuk(x) for x in satirlar[0]]
    k = {alan: next((i for i, x in enumerate(b) if x in adlar), None) for alan, adlar in ALANLAR.items()}
    if k["no"] is None or k["tutar"] is None:
        raise SystemExit(f"{yol.name}: 'Fatura No' ve 'Tutar' sütunları bulunamadı. Başlıklar: {satirlar[0]}")
    al = lambda r, a: r[k[a]] if k[a] is not None and k[a] < len(r) else None  # noqa: E731
    belgeler = []
    for n, r in enumerate(satirlar[1:], start=2):
        no = str(al(r, "no") or "").strip()
        if not no:
            continue
        belgeler.append(Belge(kaynak, n, no, tarih(al(r, "tarih")), re.sub(r"\D", "", str(al(r, "vkn") or "")),
                              str(al(r, "unvan") or "").strip(), abs(para(al(r, "tutar"))),
                              para(al(r, "kdv")) if k["kdv"] is not None else None, str(al(r, "aciklama") or "")))
    return belgeler


def efatura_belgeleri(yol: Path) -> list[Belge]:
    if yol.is_file() and yol.suffix.lower() in {".xlsx", ".xlsm", ".csv", ".txt"}:
        return liste_belgeleri(yol, "efatura")
    belgeler, gorulen = [], set()
    for n, (ad, veri) in enumerate(efatura.dosyalari_topla(yol), start=1):
        f = efatura.fatura_oku(veri, ad)
        anahtar = f.ettn or (f.no, f.satici_vkn)
        if not f.no or anahtar in gorulen:      # okunamayan veya aynı faturanın ikinci kopyası
            continue
        gorulen.add(anahtar)
        kur = f.kur if f.para_birimi != "TRY" and f.kur else Decimal(1)
        belgeler.append(Belge("efatura", n, f.no, tarih(f.tarih), f.satici_vkn, f.satici_unvan,
                              (Decimal(f.vergi_dahil) * kur).quantize(Decimal("0.01")), f.toplam_kdv * kur,
                              ad + (f" ({f.para_birimi} × {kur})" if kur != 1 else "")))
    return belgeler


@dataclass
class Eslesme:
    e: Belge
    k: Belge
    yontem: str

    @property
    def tutar_farki(self) -> Decimal:
        return self.k.tutar - self.e.tutar


def eslestir(efaturalar: list[Belge], kayitlar: list[Belge], gun_toleransi: int = 3, tutar_toleransi=Decimal("0.05")):
    acik_k = list(kayitlar)
    eslesmeler, acik_e = [], []
    no_vkn = {}
    for k in acik_k:
        no_vkn.setdefault((k.anahtar_no, k.vkn), []).append(k)
    for e in efaturalar:
        aday = None
        for anahtar, yontem in (((e.anahtar_no, e.vkn), "Fatura no + VKN"), ((e.anahtar_no, ""), "Fatura no")):
            liste = [x for x in no_vkn.get(anahtar, []) if x in acik_k]
            if liste:
                aday = (liste[0], yontem)
                break
        if aday:
            eslesmeler.append(Eslesme(e, aday[0], aday[1]))
            acik_k.remove(aday[0])
        else:
            acik_e.append(e)
    kalan_e = []
    for e in acik_e:   # 3. adım: VKN + tutar + tarih
        adaylar = [k for k in acik_k if k.vkn and k.vkn == e.vkn and abs(k.tutar - e.tutar) <= tutar_toleransi
                   and (not k.tarih or not e.tarih or abs((k.tarih - e.tarih).days) <= gun_toleransi)]
        if adaylar:
            k = min(adaylar, key=lambda x: abs((x.tarih - e.tarih).days) if x.tarih and e.tarih else 0)
            eslesmeler.append(Eslesme(e, k, "VKN + tutar + tarih (fatura no farklı)"))
            acik_k.remove(k)
        else:
            kalan_e.append(e)
    return eslesmeler, kalan_e, acik_k


BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
SARI = PatternFill("solid", fgColor="FFF4CE")
PARA = "#,##0.00"


def _sayfa(wb, ad, basliklar, satirlar, para_sutunlari=()):
    ws = wb.create_sheet(ad)
    ws.append(basliklar)
    for h in ws[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for s in satirlar:
        ws.append([float(x) if isinstance(x, Decimal) else x for x in s])
    for j, b in enumerate(basliklar, 1):
        ws.column_dimensions[get_column_letter(j)].width = max(12, min(45, len(b) + 4))
        if j in para_sutunlari:
            for h in ws[get_column_letter(j)][1:]:
                h.number_format = PARA
        for h in ws[get_column_letter(j)][1:]:
            if isinstance(h.value, date):
                h.number_format = "DD.MM.YYYY"
    ws.freeze_panes = "A2"
    if satirlar:
        ws.auto_filter.ref = ws.dimensions
    return ws


def calistir(efatura_yolu: Path, kayit_yolu: Path, cikti: Path, gun_toleransi: int = 3) -> dict:
    efaturalar = efatura_belgeleri(efatura_yolu)
    kayitlar = liste_belgeleri(kayit_yolu, "kayit")
    eslesmeler, defterde_yok, portalda_yok = eslestir(efaturalar, kayitlar, gun_toleransi)

    tutar_farkli = [m for m in eslesmeler if abs(m.tutar_farki) > Decimal("0.05")]
    sayac = Counter((k.anahtar_no, k.vkn) for k in kayitlar)
    mukerrer = [k for k in kayitlar if sayac[(k.anahtar_no, k.vkn)] > 1]

    wb = Workbook()
    wb.remove(wb.active)
    ozet = [
        ["Gelen e-fatura/e-arşiv", len(efaturalar)], ["Muhasebe alış kaydı", len(kayitlar)],
        ["Eşleşen", len(eslesmeler)], ["Deftere işlenmemiş (portalda var, kayıtta yok)", len(defterde_yok)],
        ["Portalda karşılığı yok (kayıtta var)", len(portalda_yok)], ["Tutarı farklı", len(tutar_farkli)],
        ["Mükerrer kayıt (aynı no + VKN birden çok kez)", len(mukerrer)],
        ["İşlenmemiş faturaların toplam tutarı", float(sum((e.tutar for e in defterde_yok), Decimal(0)))],
    ]
    ws = _sayfa(wb, "Özet", ["Gösterge", "Değer"], ozet)
    ws.column_dimensions["A"].width = 48
    ws.cell(9, 2).number_format = PARA
    _sayfa(wb, "Deftere İşlenmemiş", ["Fatura No", "Tarih", "Satıcı VKN", "Satıcı", "Vergiler Dahil", "KDV", "Kaynak"],
           [[e.no, e.tarih, e.vkn, e.unvan, e.tutar, e.kdv, e.aciklama] for e in sorted(defterde_yok, key=lambda x: (x.tarih or date.min))],
           para_sutunlari=(5, 6))
    _sayfa(wb, "Tutarı Farklı", ["Fatura No", "Satıcı", "e-Fatura Tutarı", "Kayıt Tutarı", "Fark", "Kayıt Satırı", "Eşleşme"],
           [[m.e.no, m.e.unvan, m.e.tutar, m.k.tutar, m.tutar_farki, m.k.sira, m.yontem] for m in tutar_farkli],
           para_sutunlari=(3, 4, 5))
    _sayfa(wb, "Portalda Yok", ["Kayıt Satırı", "Belge No", "Tarih", "VKN", "Cari", "Tutar", "Açıklama", "Olası neden"],
           [[k.sira, k.no, k.tarih, k.vkn, k.unvan, k.tutar, k.aciklama,
             "Kâğıt fatura, serbest meslek makbuzu, ithalat, yanlış dönem veya hatalı numara olabilir"] for k in portalda_yok],
           para_sutunlari=(6,))
    _sayfa(wb, "Mükerrer Kayıtlar", ["Kayıt Satırı", "Belge No", "VKN", "Cari", "Tutar", "Tarih"],
           [[k.sira, k.no, k.vkn, k.unvan, k.tutar, k.tarih] for k in sorted(mukerrer, key=lambda x: x.anahtar_no)],
           para_sutunlari=(5,))
    es = _sayfa(wb, "Eşleşenler", ["Fatura No", "Satıcı", "e-Fatura Tutarı", "Kayıt Belge No", "Kayıt Tutarı", "Eşleşme Yöntemi"],
                [[m.e.no, m.e.unvan, m.e.tutar, m.k.no, m.k.tutar, m.yontem] for m in eslesmeler], para_sutunlari=(3, 5))
    for r, m in enumerate(eslesmeler, start=2):
        if m.yontem.startswith("VKN"):
            for h in es[r]:
                h.fill = SARI
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    return {"eslesen": eslesmeler, "defterde_yok": defterde_yok, "portalda_yok": portalda_yok,
            "tutar_farkli": tutar_farkli, "mukerrer": mukerrer}


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description="Gelen e-faturaların muhasebeye işlenip işlenmediğini kontrol eder.")
    ap.add_argument("--efatura", type=Path, default=BURASI / "ornek_veri" / "gelen_efaturalar", help="XML/ZIP klasörü veya portal listesi (.xlsx/.csv)")
    ap.add_argument("--kayitlar", type=Path, default=BURASI / "ornek_veri" / "alis_kayitlari.csv", help="Muhasebe alış kayıtları (.xlsx/.csv)")
    ap.add_argument("--gun-toleransi", type=int, default=3)
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "efatura_kayit_kontrolu.xlsx")
    a = ap.parse_args(argv)
    s = calistir(a.efatura, a.kayitlar, a.cikti, a.gun_toleransi)
    print(f"[OK] Eşleşen {len(s['eslesen'])} · deftere işlenmemiş {len(s['defterde_yok'])} · portalda yok {len(s['portalda_yok'])} · "
          f"tutarı farklı {len(s['tutar_farkli'])} · mükerrer {len(s['mukerrer'])}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
