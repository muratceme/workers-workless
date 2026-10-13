"""
Nakliye Maliyet Analizi — Workers / Workless kod bloğu
Lojistik ve Depo › Lojistik Uzmanı

Nakliye faturalarını sefer, ağırlık ve bölge bazında analiz eder; birim taşıma maliyetini hesaplar.
  Kg başı maliyet       = fatura tutarı ÷ taşınan ağırlık (kg)
  Km başı maliyet       = fatura tutarı ÷ mesafe (km)
  Ton-km başı maliyet   = fatura tutarı ÷ (ağırlık ÷ 1.000 × km)
  Doluluk %             = taşınan ağırlık ÷ araç kapasitesi × 100
  Grup ortalamaları     = toplam tutar ÷ toplam ağırlık (ağırlıklı ortalama; seferlerin basit ortalaması değil)
  Bölge ortalamasına fark % = (taşıyıcının o bölgedeki kg maliyeti ÷ bölgenin kg maliyeti − 1) × 100
  Pahalı sefer          = kg maliyeti, bölgenin medyan kg maliyetinin %30 üstünde (bölgede en az 3 sefer varsa)
Tutarlar KDV hariç ve tek para biriminde olmalıdır. Aynı sefer numarası ikinci kez gelirse sayılmaz ve uyarılır.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
SUTUN = {"tarih": ("tarih", "sefer tarihi", "fatura tarihi"), "sefer": ("sefer no", "irsaliye no", "fatura no", "no"), "tasiyici": ("tasiyici", "nakliyeci", "nakliye firmasi", "firma"), "bolge": ("bolge", "varis", "varis ili", "guzergah"),
         "agirlik": ("agirlik kg", "agirlik", "kg", "tonaj kg"), "km": ("km", "mesafe", "mesafe km"), "kapasite": ("arac kapasitesi kg", "kapasite kg", "arac kapasitesi", "kapasite"), "tutar": ("fatura tutari", "tutar", "navlun", "nakliye bedeli")}
KIRILIM = (("Bölgeler", "bolge", "Bölge"), ("Taşıyıcılar", "tasiyici", "Taşıyıcı"), ("Aylar", "ay", "Ay"))
Z, YUZ, BIN = Decimal(0), Decimal(100), Decimal(1000)


def olcu(x) -> Decimal | None:
    """Ağırlık ve mesafe okur: virgül ya da tek nokta ondalık ayırıcıdır; '1.250,5' binlik ayırıcılıdır."""
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x or ""))
    try:
        return Decimal(s.replace(".", "").replace(",", ".") if "," in s else s) if s else None
    except InvalidOperation:
        return None


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("agirlik", "tutar")):
        if not metin(r["agirlik"]) and not metin(r["tutar"]):
            continue
        kim = metin(r["sefer"]) or f"Satır {r['_satir']}"
        agirlik, tutar, km, kapasite = olcu(r["agirlik"]), para(r["tutar"]), olcu(r["km"]), olcu(r["kapasite"])
        if agirlik is None or agirlik <= 0 or tutar is None or tutar < 0 or any(metin(r[a]) and (v is None or v <= 0) for a, v in (("km", km), ("kapasite", kapasite))):
            uy.append(uyari("Yüksek", "Sefer okunamadı", kim, f"Satır {r['_satir']}: ağırlık '{metin(r['agirlik'])}', tutar '{metin(r['tutar'])}', km '{metin(r['km'])}', kapasite '{metin(r['kapasite'])}'. Ağırlık, km ve kapasite pozitif olmalı. "
                            "Satır alınmadı."))
            continue
        t = tarih(r["tarih"])
        l.append({"tarih": t, "ay": f"{t:%Y-%m}" if t else "(tarih yok)", "sefer": metin(r["sefer"]), "tasiyici": metin(r["tasiyici"]) or "(taşıyıcı yok)", "bolge": metin(r["bolge"]) or "(bölge yok)", "agirlik": agirlik, "km": km,
                  "kapasite": kapasite, "tutar": tutar})
    return l, uy


def _bol(pay: Decimal, payda: Decimal) -> Decimal | None:
    return yuvarla(pay / payda) if payda else None


def _topla(liste: list[dict]) -> dict:
    kmli, kapasiteli = [s for s in liste if s["km"]], [s for s in liste if s["kapasite"]]
    agirlik, tutar = sum((s["agirlik"] for s in liste), Z), sum((s["tutar"] for s in liste), Z)
    return {"sefer": len(liste), "agirlik": agirlik, "km": sum((s["km"] for s in kmli), Z) if kmli else None, "tutar": tutar, "kg_maliyet": _bol(tutar, agirlik), "km_maliyet": _bol(sum((s["tutar"] for s in kmli), Z), sum((s["km"] for s in kmli), Z)),
            "tonkm_maliyet": _bol(sum((s["tutar"] for s in kmli), Z), sum((s["agirlik"] / BIN * s["km"] for s in kmli), Z)), "sefer_basi": _bol(tutar, Decimal(len(liste))),
            "doluluk": _bol(sum((s["agirlik"] for s in kapasiteli), Z) * YUZ, sum((s["kapasite"] for s in kapasiteli), Z))}


def _medyan(d: list[Decimal]) -> Decimal:
    d = sorted(d)
    return d[len(d) // 2] if len(d) % 2 else (d[len(d) // 2 - 1] + d[len(d) // 2]) / 2


def analiz(seferler: list[dict], sapma: Decimal = Decimal(30), doluluk_esik: Decimal = Decimal(70), asgari: int = 3) -> dict:
    uy, liste, gorulen = [], [], set()
    for s in seferler:
        if s["sefer"] and katla(s["sefer"]) in gorulen:
            uy.append(uyari("Yüksek", "Mükerrer sefer numarası", s["sefer"], f"Aynı sefer / fatura numarası ikinci kez geçiyor ({s['tasiyici']}, {s['tutar']:,.2f}). İkinci satır hesaba katılmadı; çift fatura olup olmadığını kontrol edin."))
            continue
        gorulen.add(katla(s["sefer"]))
        liste.append(s)
    if not liste:
        raise ValueError("Analiz edilecek sefer yok.")
    for s in liste:
        s["kg_maliyet"], s["km_maliyet"] = s["tutar"] / s["agirlik"], (s["tutar"] / s["km"] if s["km"] else None)
        s["tonkm_maliyet"] = s["tutar"] / (s["agirlik"] / BIN * s["km"]) if s["km"] else None
        s["doluluk"] = yuvarla(s["agirlik"] * YUZ / s["kapasite"]) if s["kapasite"] else None
    bolgeler = defaultdict(list)
    for s in liste:
        bolgeler[katla(s["bolge"])].append(s)
    for s in liste:
        kim = s["sefer"] or f"{s['tarih']:%d.%m.%Y} {s['bolge']}" if s["sefer"] or s["tarih"] else s["bolge"]
        grup = bolgeler[katla(s["bolge"])]
        s["medyan"] = yuvarla(_medyan([x["kg_maliyet"] for x in grup])) if len(grup) >= asgari else None
        s["pahali"] = len(grup) >= asgari and s["kg_maliyet"] > _medyan([x["kg_maliyet"] for x in grup]) * (1 + sapma / YUZ)
        if s["pahali"]:
            uy.append(uyari("Orta", "Bölge medyanının üstünde sefer", kim, f"{s['bolge']} · {s['tasiyici']}: kg başı {yuvarla(s['kg_maliyet'])}, bölge medyanı {s['medyan']} (eşik +%{sapma})."
                            + (f" Doluluk %{s['doluluk']}." if s["doluluk"] is not None else "")))
        if s["doluluk"] is not None and s["doluluk"] > 100:
            uy.append(uyari("Orta", "Kapasite aşımı", kim, f"Taşınan {s['agirlik']:f} kg, araç kapasitesi {s['kapasite']:f} kg (%{s['doluluk']}). Ağırlık ya da kapasite bilgisini kontrol edin."))
        elif s["doluluk"] is not None and s["doluluk"] < doluluk_esik:
            uy.append(uyari("Bilgi", "Düşük doluluk", kim, f"{s['bolge']} · {s['tasiyici']}: doluluk %{s['doluluk']} (eşik %{doluluk_esik}); kg başı maliyet {yuvarla(s['kg_maliyet'])}."))
    kirilimlar = {}
    for baslik, alan, _ in KIRILIM:
        grup = {}
        for s in liste:
            grup.setdefault(s[alan] if alan == "ay" else katla(s[alan]), []).append(s)
        kirilimlar[baslik] = [{"ad": g[0][alan], **_topla(g)} for _, g in sorted(grup.items())]
    bolge_kg = {katla(b["ad"]): b["tutar"] / b["agirlik"] for b in kirilimlar["Bölgeler"]}
    ciftler, yazim = {}, {}
    for s in liste:
        ciftler.setdefault((katla(s["bolge"]), katla(s["tasiyici"])), []).append(s)
        yazim.setdefault(katla(s["bolge"]), s["bolge"])
        yazim.setdefault(katla(s["tasiyici"]), s["tasiyici"])
    eslesme = []
    for (b, tk), g in sorted(ciftler.items()):
        t = _topla(g)
        eslesme.append({"bolge": yazim[b], "tasiyici": yazim[tk], "sefer": t["sefer"], "agirlik": t["agirlik"], "tutar": t["tutar"], "kg_maliyet": t["kg_maliyet"],
                        "fark": yuvarla((t["tutar"] / t["agirlik"] / bolge_kg[b] - 1) * YUZ) if bolge_kg[b] else None})
    for s in liste:
        for alan in ("kg_maliyet", "km_maliyet", "tonkm_maliyet"):
            s[alan] = yuvarla(s[alan]) if s[alan] is not None else None
    liste.sort(key=lambda s: (s["tarih"] or date.min, katla(s["sefer"])))
    return {"seferler": liste, "toplam": _topla(liste), "kirilimlar": kirilimlar, "eslesme": eslesme, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    olcut = ["Sefer", "Ağırlık (kg)", "Km", "Tutar", "Kg Başı", "Km Başı", "Ton-km Başı", "Sefer Başı Tutar", "Doluluk %"]
    deger = lambda o: [o["sefer"], o["agirlik"], o["km"], o["tutar"], o["kg_maliyet"], o["km_maliyet"], o["tonkm_maliyet"], o["sefer_basi"], o["doluluk"]]   # noqa: E731
    for i, (baslik, _, ad) in enumerate(KIRILIM):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = baslik
        _baslik(ws, [ad] + olcut, (26, 8, 14, 11, 16, 10, 10, 11, 15, 10))
        for o in r["kirilimlar"][baslik] + [{"ad": "TOPLAM", **r["toplam"]}]:
            n = _satir(ws, [o["ad"]] + deger(o), para_s=(3, 4, 5, 6, 7, 8, 9, 10))
        for h in ws[n]:
            h.font = Font(bold=True)
    ws = wb.create_sheet("Bölge-Taşıyıcı")
    _baslik(ws, ["Bölge", "Taşıyıcı", "Sefer", "Ağırlık (kg)", "Tutar", "Kg Başı", "Bölge Ortalamasına Fark %"], (22, 26, 8, 14, 16, 10, 14))
    for e in r["eslesme"]:
        _satir(ws, [e["bolge"], e["tasiyici"], e["sefer"], e["agirlik"], e["tutar"], e["kg_maliyet"], e["fark"]], para_s=(4, 5, 6, 7), dolgu={7: None if e["fark"] is None else YESIL if e["fark"] < 0 else SARI if e["fark"] > 0 else None})
    ws = wb.create_sheet("Seferler")
    _baslik(ws, ["Tarih", "Sefer No", "Taşıyıcı", "Bölge", "Ağırlık (kg)", "Km", "Kapasite (kg)", "Tutar", "Kg Başı", "Km Başı", "Ton-km Başı", "Doluluk %", "Bölge Medyanı (kg başı)"], (12, 14, 24, 20, 13, 9, 13, 15, 10, 10, 11, 10, 13))
    for s in r["seferler"]:
        _satir(ws, [s["tarih"], s["sefer"], s["tasiyici"], s["bolge"], s["agirlik"], s["km"], s["kapasite"], s["tutar"], s["kg_maliyet"], s["km_maliyet"], s["tonkm_maliyet"], s["doluluk"], s["medyan"]], para_s=(5, 6, 7, 8, 9, 10, 11, 12, 13),
               tarih_s=(1,), dolgu={9: SARI if s["pahali"] else None, 12: KIRMIZI if s["doluluk"] is not None and s["doluluk"] > 100 else None})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, sapma: Decimal = Decimal(30), doluluk_esik: Decimal = Decimal(70), asgari: int = 3) -> dict:
    seferler, uy = oku(yol)
    r = analiz(seferler, sapma, doluluk_esik, asgari)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Nakliye faturalarını sefer, ağırlık ve bölge bazında analiz edip birim taşıma maliyetini hesaplar.")
    p.add_argument("--seferler", type=Path, help="Tarih, Sefer No, Taşıyıcı, Bölge, Ağırlık (kg), Km, Araç Kapasitesi (kg), Fatura Tutarı")
    p.add_argument("--sapma", default="30", help="Kg maliyeti bölge medyanını bu yüzdeden fazla aşan sefer işaretlenir (varsayılan 30)")
    p.add_argument("--doluluk", default="70", help="Doluluğu bu yüzdenin altındaki seferler bildirilir (varsayılan 70)")
    p.add_argument("--asgari", type=int, default=3, help="Medyan karşılaştırması için bölgede gereken en az sefer (varsayılan 3)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "nakliye_maliyet_analizi.xlsx")
    a = p.parse_args(argv)
    yol = a.seferler or ORNEK / "seferler.csv"
    sapma, doluluk = para(a.sapma), para(a.doluluk)
    if sapma is None or sapma <= 0 or doluluk is None or not 0 <= doluluk <= 100 or a.asgari < 2:
        print("[X] --sapma pozitif; --doluluk 0–100 arasında; --asgari en az 2 olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, sapma, doluluk, a.asgari)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {t['sefer']} sefer · {t['agirlik']:,.0f} kg · tutar {t['tutar']:,.2f} · kg başı {t['kg_maliyet']}" + (f" · km başı {t['km_maliyet']}" if t["km_maliyet"] is not None else "") + (f" · doluluk %{t['doluluk']}" if t["doluluk"] is not None else ""))
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
