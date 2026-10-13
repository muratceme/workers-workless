import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def s(sefer, agirlik, tutar, bolge="B", tasiyici="T", km=None, kapasite=None, t=date(2026, 9, 1)):
    d = lambda x: None if x is None else D(str(x))   # noqa: E731
    return {"tarih": t, "ay": f"{t:%Y-%m}" if t else "(tarih yok)", "sefer": sefer, "tasiyici": tasiyici, "bolge": bolge, "agirlik": d(agirlik), "km": d(km), "kapasite": d(kapasite), "tutar": d(tutar)}


class AnalizTesti(unittest.TestCase):
    def test_birim_maliyet(self):
        r = main.analiz([s("1", 8000, 20000, km=450, kapasite=20000), s("2", 3000, 1000), s("3", 1000, 500, km=100)])
        self.assertEqual([(x["sefer"], x["kg_maliyet"], x["km_maliyet"], x["tonkm_maliyet"], x["doluluk"]) for x in r["seferler"]],
                         [("1", D("2.50"), D("44.44"), D("5.56"), D("40.00")), ("2", D("0.33"), None, None, None), ("3", D("0.50"), D("5.00"), D("5.00"), None)])
        self.assertEqual(r["toplam"], {"sefer": 3, "agirlik": D(12000), "km": D(550), "tutar": D(21500), "kg_maliyet": D("1.79"), "km_maliyet": D("37.27"), "tonkm_maliyet": D("5.54"), "sefer_basi": D("7166.67"), "doluluk": D("40.00")})
        t = main.analiz([s("1", 100, 50)])["toplam"]
        self.assertEqual((t["km"], t["km_maliyet"], t["tonkm_maliyet"], t["doluluk"]), (None, None, None, None))
        with self.assertRaises(ValueError):
            main.analiz([])

    def test_medyan_ve_doluluk(self):
        p = lambda tutarlar, **kw: [(x["pahali"], x["medyan"]) for x in main.analiz([s(str(i), 100, t) for i, t in enumerate(tutarlar)], **kw)["seferler"]]   # noqa: E731
        self.assertEqual(p((100, 100, 130)), [(False, D("1.00"))] * 3)
        self.assertEqual(p((100, 100, 131)), [(False, D("1.00")), (False, D("1.00")), (True, D("1.00"))])
        self.assertEqual(p((100, 200, 300, 1000)), [(False, D("2.50"))] * 3 + [(True, D("2.50"))])
        self.assertEqual((p((100, 900)), p((100, 900), asgari=2), p((100, 100, 121), sapma=D(20))[2]), ([(False, None)] * 2, [(False, D("5.00")), (True, D("5.00"))], (True, D("1.00"))))
        r = main.analiz([s("a", 70, 1, kapasite=100), s("b", "69.99", 1, kapasite=100), s("c", 100, 1, kapasite=100), s("d", "100.01", 1, kapasite=100), s("e", 80, 1, kapasite=100)], doluluk_esik=D(70))
        self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]], [("Bilgi", "Düşük doluluk", "b"), ("Orta", "Kapasite aşımı", "d")])
        self.assertEqual([u["kim"] for u in main.analiz([s("e", 80, 1, kapasite=100)], doluluk_esik=D(85))["uyarilar"]], ["e"])
        r = main.analiz([s("", 100, 500, "İzmir", t=None), s("", 100, 100, "izmir"), s("", 100, 100, "İZMİR")])
        self.assertEqual([u["kim"] for u in r["uyarilar"]], ["İzmir"])

    def test_kirilim_ve_mukerrer(self):
        l = [s("1", 1000, 1500, "Ankara", "X", t=date(2026, 9, 3)), s("2", 3000, 3000, "ankara", "Y", t=date(2026, 10, 1)), s("3", 500, 1000, "Bursa", "x", t=date(2026, 9, 9)), s("1", 9, 9, "Bursa", "X"), s("", 500, 500, "Bursa", "Y", t=None),
             s("", 100, 100, "Bursa", "Y", t=None)]
        r = main.analiz(l)
        k = lambda ad: [(o["ad"], o["sefer"], o["agirlik"], o["tutar"], o["kg_maliyet"], o["sefer_basi"]) for o in r["kirilimlar"][ad]]   # noqa: E731
        self.assertEqual(k("Bölgeler"), [("Ankara", 2, D(4000), D(4500), D("1.13"), D("2250.00")), ("Bursa", 3, D(1100), D(1600), D("1.45"), D("533.33"))])
        self.assertEqual(k("Taşıyıcılar"), [("X", 2, D(1500), D(2500), D("1.67"), D("1250.00")), ("Y", 3, D(3600), D(3600), D("1.00"), D("1200.00"))])
        self.assertEqual(k("Aylar"), [("(tarih yok)", 2, D(600), D(600), D("1.00"), D("300.00")), ("2026-09", 2, D(1500), D(2500), D("1.67"), D("1250.00")), ("2026-10", 1, D(3000), D(3000), D("1.00"), D("3000.00"))])
        self.assertEqual([(e["bolge"], e["tasiyici"], e["sefer"], e["kg_maliyet"], e["fark"]) for e in r["eslesme"]],
                         [("Ankara", "X", 1, D("1.50"), D("33.33")), ("Ankara", "Y", 1, D("1.00"), D("-11.11")), ("Bursa", "X", 1, D("2.00"), D("37.50")), ("Bursa", "Y", 2, D("1.00"), D("-31.25"))])
        self.assertEqual(([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]], r["toplam"]["sefer"], [x["sefer"] for x in r["seferler"]]), ([("Yüksek", "Mükerrer sefer numarası", "1"), ("Orta", "Bölge medyanının üstünde sefer", "3")], 5, ["", "", "1", "3", "2"]))


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "seferler.csv", cikti)
            self.assertEqual(r["toplam"], {"sefer": 11, "agirlik": D(139400), "km": D(4140), "tutar": D(220690), "kg_maliyet": D("1.58"), "km_maliyet": D("53.31"), "tonkm_maliyet": D("3.80"), "sefer_basi": D("20062.73"), "doluluk": D("83.98")})
            k = lambda ad: [(o["ad"], o["sefer"], o["agirlik"], o["tutar"], o["kg_maliyet"], o["km_maliyet"], o["tonkm_maliyet"], o["doluluk"]) for o in r["kirilimlar"][ad]]   # noqa: E731
            self.assertEqual(k("Bölgeler"), [("Ankara", 5, D(81000), D(129100), D("1.59"), D("57.38"), D("3.54"), D("81.00")), ("Bursa", 3, D(19400), D(23440), D("1.21"), D("52.09"), D("8.05"), D("80.83")),
                                             ("İzmir", 3, D(39000), D(68150), D("1.75"), D("47.33"), D("3.64"), D("92.86"))])
            self.assertEqual([(o["ad"], o["sefer"], o["agirlik"], o["tutar"], o["kg_maliyet"]) for o in r["kirilimlar"]["Taşıyıcılar"]],
                             [("Ege Taşıma", 3, D(35400), D(55790), D("1.58")), ("Hızlı Nakliyat", 4, D(69000), D(105100), D("1.52")), ("Yıldız Lojistik", 4, D(35000), D(59800), D("1.71"))])
            self.assertEqual([(o["ad"], o["sefer"], o["tutar"], o["kg_maliyet"]) for o in r["kirilimlar"]["Aylar"]], [("2026-09", 10, D(192190), D("1.60")), ("2026-10", 1, D(28500), D("1.50"))])
            self.assertEqual([(e["bolge"], e["tasiyici"], e["sefer"], e["kg_maliyet"], e["fark"]) for e in r["eslesme"]],
                             [("Ankara", "Hızlı Nakliyat", 3, D("1.46"), D("-8.09")), ("Ankara", "Yıldız Lojistik", 2, D("1.90"), D("19.21")), ("Bursa", "Ege Taşıma", 1, D("1.10"), D("-8.96")), ("Bursa", "Yıldız Lojistik", 2, D("1.29"), D("6.84")),
                              ("İzmir", "Ege Taşıma", 2, D("1.72"), D("-1.34")), ("İzmir", "Hızlı Nakliyat", 1, D("1.80"), D("3.01"))])
            x = {v["sefer"]: v for v in r["seferler"]}
            self.assertEqual([(x[n]["kg_maliyet"], x[n]["doluluk"], x[n]["medyan"], x[n]["pahali"]) for n in ("SF-004", "SF-009", "SF-010")],
                             [(D("2.50"), D("40.00"), D("1.50"), True), (D("1.40"), D("62.50"), D("1.20"), False), (D("1.10"), D("105.00"), D("1.20"), False)])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Yüksek", "Mükerrer sefer numarası", "SF-011"), ("Orta", "Bölge medyanının üstünde sefer", "SF-004"), ("Orta", "Kapasite aşımı", "SF-010"), ("Bilgi", "Düşük doluluk", "SF-004"), ("Bilgi", "Düşük doluluk", "SF-009")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Bölgeler", "Taşıyıcılar", "Aylar", "Bölge-Taşıyıcı", "Seferler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "s.csv"
            y.write_text("Fatura Tarihi;Fatura No;Nakliyeci;Varış;Kg;Mesafe;Kapasite;Navlun\n05.09.2026;F1;;;1.250,5;450;;12.500,75\nx;F2;A;B;12.5;;2.5;100\n;F3;A;B;0;;;5\n;F4;A;B;5;;;x\n;F5;A;B;5;0;;5\n;;;;;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["tarih"], x["ay"], x["sefer"], x["tasiyici"], x["bolge"], x["agirlik"], x["km"], x["kapasite"], x["tutar"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([(date(2026, 9, 5), "2026-09", "F1", "(taşıyıcı yok)", "(bölge yok)", D("1250.5"), D(450), None, D("12500.75")), (None, "(tarih yok)", "F2", "A", "B", D("12.5"), None, D("2.5"), D(100))],
                              [("Sefer okunamadı", "F3"), ("Sefer okunamadı", "F4"), ("Sefer okunamadı", "F5")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--seferler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--sapma", "20", "--doluluk", "80", "--asgari", "5"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + x) for x in (["--sapma", "0"], ["--sapma", "x"], ["--doluluk", "101"], ["--asgari", "1"])], [2] * 4)
            self.assertEqual(main.main(["--seferler", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Tarih;Tutar\n05.09.2026;1\n", "Ağırlık;Tutar\n;\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
