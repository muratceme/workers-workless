import sys
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))

import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

FIX = Path(__file__).resolve().parent / "fixtures"


class GibOrnegiTesti(unittest.TestCase):
    """GİB'in resmî tevkifatlı örnek faturası (SRM2012000086985)."""

    def setUp(self):
        self.f = main.fatura_oku((FIX / "gib_ornek_tevkifat.xml").read_bytes(), "gib.xml")

    def test_baslik(self):
        f = self.f
        self.assertEqual((f.no, f.senaryo, f.tip, f.tarih), ("SRM2012000086985", "TICARIFATURA", "TEVKIFAT", "2012-08-04"))
        self.assertEqual(f.satici_vkn, "0890890895")
        self.assertEqual(f.alici_vkn, "1870200192")

    def test_tutarlar(self):
        f = self.f
        self.assertEqual(f.kdv, {Decimal("18"): [Decimal("20000"), Decimal("3600")]})
        self.assertEqual(f.tevkifat, Decimal("3240"))
        self.assertEqual(f.tevkifat_kodlari, ["606 (%90)"])
        self.assertEqual(f.odenecek, Decimal("20360"))
        self.assertEqual(len(f.satirlar), 1)
        self.assertEqual(f.satirlar[0].birim, "Adet")
        # tutarlar tutarlı: vergiler dahil − tevkifat = ödenecek
        self.assertFalse([u for u in f.uyarilar if "ödenecek" in u or "KDV:" in u])


class OrnekVeriTesti(unittest.TestCase):
    def test_uctan_uca(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "f.xlsx"
            faturalar = main.calistir(KLASOR / "ornek_veri" / "faturalar", cikti)
            self.assertEqual(len(faturalar), 8)   # 5 xml + zip içinde 2 + bozuk 1
            by = {f.no: f for f in faturalar if f.no}

            tev = by["ORN2026000000102"]
            self.assertEqual(tev.tevkifat, Decimal("800.00"))          # 4000 KDV × %20 (2/10)
            self.assertEqual(tev.odenecek, Decimal("23200.00"))

            ear = by["EAR2026000000031"]
            self.assertEqual((ear.alici_vkn, ear.alici_unvan), ("17291716060", "Ayşe Yılmaz"))

            ihr = by["IHR2026000000007"]
            self.assertEqual((ihr.para_birimi, ihr.kur), ("USD", Decimal("33.8750")))

            self.assertTrue(any("KDV:" in u for u in by["ORN2026000000110"].uyarilar))
            self.assertTrue(any("okunamadı" in u for f in faturalar for u in f.uyarilar))

            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Faturalar", "Satırlar", "KDV Özeti", "Kontroller"])
            kontroller = [r[2] for r in wb["Kontroller"].iter_rows(min_row=2, values_only=True)]
            self.assertEqual(sum("Mükerrer ETTN" in k for k in kontroller), 2)
            # KDV özeti mükerrer faturayı iki kez saymaz: %1 matrahı = 34788 + 1280 + 18600
            ozet = {(r[0], r[1]): r[2] for r in wb["KDV Özeti"].iter_rows(min_row=2, values_only=True)}
            self.assertAlmostEqual(ozet[("TRY", 1)], 34788 + 1280 + 18600, places=2)


if __name__ == "__main__":
    unittest.main()
