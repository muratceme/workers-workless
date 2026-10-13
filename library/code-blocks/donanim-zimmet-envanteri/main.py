"""
Donanım Zimmet Envanteri — Workers / Workless kod bloğu
Bilgi Teknolojileri › BT Destek Uzmanı

Bilgisayar, telefon ve çevre birimlerinin zimmet, garanti ve yenileme durumunu tek envanterde denetler.
  Garanti durumu      = garanti bitişi --tarih gününden önceyse "Bitmiş", --gun (60) gün içindeyse "Bitiyor", değilse "Devam ediyor"
  Yenileme tarihi     = alım tarihi + yenileme ömrü (satırdaki "Yenileme Ömrü (yıl)"; boşsa --omur, varsayılan 4 yıl)
  Yaş (yıl)           = (--tarih − alım tarihi) ÷ 365,25
  Zimmet denetimleri  = kullanımda ama zimmetlisi yok · depoda/hurda ama zimmet düşülmemiş · aynı demirbaş veya seri numarası
                        iki kayıtta · zimmetli kişi --personel listesinde yok · bir kişide aynı türden birden çok cihaz
Envanter sizin kaydınızdır; kod cihazın fiilen kimde olduğunu bilemez. Sayımla doğrulayın.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
SUTUN = {"no": ("demirbas no", "envanter no", "barkod", "no"), "tur": ("tur", "cihaz turu", "kategori"), "model": ("marka model", "model", "marka", "cihaz"), "seri": ("seri no", "seri numarasi", "imei"),
         "kisi": ("zimmetli", "zimmetli kisi", "kullanici", "personel"), "departman": ("departman", "birim", "bolum"), "zimmet": ("zimmet tarihi", "teslim tarihi"), "alim": ("alim tarihi", "satin alma tarihi", "fatura tarihi"),
         "garanti": ("garanti bitis", "garanti bitis tarihi", "garanti"), "durum": ("durum",), "omur": ("yenileme omru", "yenileme omru yil", "omur", "omur yil")}
PERSONEL = {"ad": ("ad soyad", "adi soyadi", "personel", "calisan", "isim")}
DURUM = {"kullanimda": "Kullanımda", "aktif": "Kullanımda", "zimmetli": "Kullanımda", "depoda": "Depoda", "stok": "Depoda", "stokta": "Depoda", "yedek": "Depoda", "bosta": "Depoda", "arizali": "Arızalı", "serviste": "Arızalı",
         "tamirde": "Arızalı", "hurda": "Hurda", "kayip": "Hurda", "satildi": "Hurda", "imha": "Hurda"}
DURUMLAR = ("Kullanımda", "Depoda", "Arızalı", "Hurda")
G_RENK = {"Bitmiş": KIRMIZI, "Bitiyor": SARI, "Devam ediyor": YESIL}


def yil_ekle(t: date, n: int) -> date:
    try:
        return t.replace(year=t.year + n)
    except ValueError:
        return t.replace(year=t.year + n, day=28)


def personel_oku(yol: Path) -> set[str]:
    return {katla(r["ad"]) for r in kayitlar(yol, PERSONEL, ("ad",)) if metin(r["ad"])}


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("no", "tur")):
        no = metin(r["no"])
        if not no and not metin(r["tur"]):
            continue
        if not no:
            uy.append(uyari("Yüksek", "Demirbaş numarası yok", f"{metin(r['tur'])} (satır {r['_satir']})", "Demirbaş numarası boş; satır envantere alınmadı."))
            continue
        kisi, d = metin(r["kisi"]), katla(r["durum"])
        if d and d not in DURUM:
            uy.append(uyari("Orta", "Durum tanınmadı", no, f"'{metin(r['durum'])}' tanınmadı; Kullanımda, Depoda, Arızalı veya Hurda yazın. Zimmet bilgisine göre varsayıldı."))
        for alan in ("alim", "garanti"):
            if metin(r[alan]) and tarih(r[alan]) is None:
                uy.append(uyari("Orta", "Tarih okunamadı", no, f"'{metin(r[alan])}' tarihi okunamadı; ilgili denetim yapılmadı."))
        omur = para(r["omur"])
        l.append({"no": no, "tur": metin(r["tur"]) or "(tür yok)", "model": metin(r["model"]), "seri": metin(r["seri"]), "kisi": kisi, "departman": metin(r["departman"]), "zimmet": tarih(r["zimmet"]), "alim": tarih(r["alim"]),
                  "garanti": tarih(r["garanti"]), "durum": DURUM.get(d, "Kullanımda" if kisi else "Depoda"), "omur": int(omur) if omur and omur > 0 else None})
    return l, uy


def denetle(cihazlar: list[dict], bugun: date, gun: int = 60, omur: int = 4, personel: set[str] | None = None) -> dict:
    if not cihazlar:
        raise ValueError("Envanterde cihaz yok.")
    uy = []
    for alan, ad in (("no", "demirbaş numarası"), ("seri", "seri numarası")):
        grup = defaultdict(list)
        for c in cihazlar:
            if c[alan]:
                grup[katla(c[alan])].append(c)
        for g in grup.values():
            if len(g) > 1:
                uy.append(uyari("Yüksek", f"Mükerrer {ad}", g[0][alan], f"Aynı {ad} {len(g)} kayıtta var: " + ", ".join(f"{c['no']} ({c['tur']}, {c['kisi'] or 'zimmetsiz'})" for c in g) + "."))
    coklu = defaultdict(list)
    for c in cihazlar:
        aktif = c["durum"] != "Hurda"
        c["yas"] = yuvarla(Decimal((bugun - c["alim"]).days) / Decimal("365.25"), "0.1") if c["alim"] else None
        c["garanti_durum"] = "" if not c["garanti"] or not aktif else "Bitmiş" if c["garanti"] < bugun else "Bitiyor" if (c["garanti"] - bugun).days <= gun else "Devam ediyor"
        c["yenileme"] = yil_ekle(c["alim"], c["omur"] or omur) if c["alim"] else None
        c["yenileme_durum"] = "" if not c["yenileme"] or not aktif else "Zamanı gelmiş" if c["yenileme"] <= bugun else "Yaklaşıyor" if (c["yenileme"] - bugun).days <= gun else "Var"
        if c["durum"] == "Kullanımda" and not c["kisi"]:
            uy.append(uyari("Yüksek", "Zimmetsiz cihaz", c["no"], f"{c['tur']} {c['model']} kullanımda görünüyor ama zimmetli kişi yazılmamış.".replace("  ", " ")))
        if c["durum"] in ("Depoda", "Hurda") and c["kisi"]:
            uy.append(uyari("Orta", "Zimmet düşülmemiş", c["no"], f"Cihaz '{c['durum']}' durumunda ama hâlâ {c['kisi']} üzerine zimmetli. Zimmet iade tutanağını ve kaydı güncelleyin."))
        if personel is not None and c["kisi"] and aktif and katla(c["kisi"]) not in personel:
            uy.append(uyari("Yüksek", "Zimmetli kişi personel listesinde yok", c["no"], f"{c['tur']} {c['kisi']} üzerine zimmetli; kişi güncel personel listesinde bulunamadı. Ayrıldıysa cihazı geri alın."))
        if c["garanti_durum"] == "Bitmiş":
            uy.append(uyari("Bilgi", "Garantisi bitmiş", c["no"], f"{c['tur']} {c['model']}: garanti {c['garanti']:%d.%m.%Y} tarihinde bitti."))
        elif c["garanti_durum"] == "Bitiyor":
            uy.append(uyari("Orta", "Garantisi bitiyor", c["no"], f"{c['tur']} {c['model']}: garanti {c['garanti']:%d.%m.%Y} tarihinde bitiyor ({(c['garanti'] - bugun).days} gün). Bilinen arıza varsa şimdi başvurun."))
        if c["durum"] == "Arızalı" and c["garanti_durum"] in ("Bitiyor", "Devam ediyor"):
            uy.append(uyari("Orta", "Arızalı cihaz garanti kapsamında", c["no"], f"{c['tur']} {c['model']} arızalı ve garantisi {c['garanti']:%d.%m.%Y} tarihine kadar sürüyor."))
        if c["yenileme_durum"] == "Zamanı gelmiş":
            uy.append(uyari("Orta", "Yenileme zamanı gelmiş", c["no"], f"{c['tur']} {c['model']}: alım {c['alim']:%d.%m.%Y}, yaş {c['yas']} yıl; planlanan yenileme {c['yenileme']:%d.%m.%Y}."))
        elif c["yenileme_durum"] == "Yaklaşıyor":
            uy.append(uyari("Bilgi", "Yenileme yaklaşıyor", c["no"], f"{c['tur']} {c['model']}: planlanan yenileme {c['yenileme']:%d.%m.%Y}."))
        if c["kisi"] and c["durum"] == "Kullanımda":
            coklu[(katla(c["kisi"]), katla(c["tur"]))].append(c)
    for g in coklu.values():
        if len(g) > 1:
            uy.append(uyari("Bilgi", "Bir kişide aynı türden birden çok cihaz", f"{g[0]['kisi']} · {g[0]['tur']}", f"{len(g)} cihaz: {', '.join(c['no'] for c in g)}. Eski cihaz iade edilmemiş olabilir."))
    kisiler, turler = {}, {}
    for c in cihazlar:
        t = turler.setdefault(katla(c["tur"]), {"tur": c["tur"], "adet": 0, **dict.fromkeys(DURUMLAR, 0), "garanti_bitmis": 0, "yenilenecek": 0})
        t["adet"] += 1
        t[c["durum"]] += 1
        t["garanti_bitmis"] += c["garanti_durum"] == "Bitmiş"
        t["yenilenecek"] += c["yenileme_durum"] == "Zamanı gelmiş"
        if c["kisi"] and c["durum"] != "Hurda":
            k = kisiler.setdefault(katla(c["kisi"]), {"kisi": c["kisi"], "departman": c["departman"], "adet": 0, "cihazlar": []})
            k["adet"] += 1
            k["departman"] = k["departman"] or c["departman"]
            k["cihazlar"].append(f"{c['tur']} ({c['no']})")
    return {"cihazlar": cihazlar, "kisiler": sorted(kisiler.values(), key=lambda k: katla(k["kisi"])), "turler": sorted(turler.values(), key=lambda t: (-t["adet"], katla(t["tur"]))), "bugun": bugun,
            "personel": personel is not None, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Envanter"
    _baslik(ws, ["Demirbaş No", "Tür", "Marka / Model", "Seri No", "Durum", "Zimmetli", "Departman", "Zimmet Tarihi", "Alım Tarihi", "Yaş (yıl)", "Garanti Bitiş", "Garanti", "Planlanan Yenileme", "Yenileme"],
            (13, 20, 20, 16, 12, 22, 16, 12, 12, 8, 12, 13, 13, 14))
    for c in r["cihazlar"]:
        _satir(ws, [c["no"], c["tur"], c["model"], c["seri"], c["durum"], c["kisi"], c["departman"], c["zimmet"], c["alim"], c["yas"], c["garanti"], c["garanti_durum"], c["yenileme"], c["yenileme_durum"]], tarih_s=(8, 9, 11, 13),
               dolgu={6: KIRMIZI if c["durum"] == "Kullanımda" and not c["kisi"] else None, 12: G_RENK.get(c["garanti_durum"]), 14: SARI if c["yenileme_durum"] == "Zamanı gelmiş" else None})
    ws = wb.create_sheet("Kişi Özeti")
    _baslik(ws, ["Zimmetli", "Departman", "Cihaz Sayısı", "Cihazlar"], (24, 18, 10, 90))
    for k in r["kisiler"]:
        _satir(ws, [k["kisi"], k["departman"], k["adet"], ", ".join(k["cihazlar"])], sar=(4,))
    ws = wb.create_sheet("Tür Özeti")
    _baslik(ws, ["Tür", "Adet"] + list(DURUMLAR) + ["Garantisi Bitmiş", "Yenileme Zamanı Gelmiş"], (24, 8, 12, 9, 9, 8, 12, 14))
    for t in r["turler"]:
        _satir(ws, [t["tur"], t["adet"]] + [t[d] or None for d in DURUMLAR] + [t["garanti_bitmis"] or None, t["yenilenecek"] or None])
    ws.append([])
    ws.append([f"Durum tarihi {r['bugun']:%d.%m.%Y}." + ("" if r["personel"] else " Personel listesi verilmedi; ayrılan çalışan denetimi yapılmadı.")])
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, gun: int = 60, omur: int = 4, personel_yol: Path | None = None) -> dict:
    cihazlar, uy = oku(yol)
    r = denetle(cihazlar, bugun, gun, omur, personel_oku(personel_yol) if personel_yol else None)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Donanım zimmet envanterini garanti, yenileme ve zimmet tutarlılığı açısından denetler.")
    p.add_argument("--envanter", type=Path, help="Demirbaş No, Tür, Marka Model, Seri No, Zimmetli, Departman, Zimmet Tarihi, Alım Tarihi, Garanti Bitiş, Durum, Yenileme Ömrü (yıl)")
    p.add_argument("--personel", type=Path, help="Güncel personel listesi (Ad Soyad); verilirse listede olmayan zimmetliler işaretlenir")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--gun", type=int, default=60, help="Garanti ve yenileme için kaç gün önceden uyarılsın (varsayılan 60)")
    p.add_argument("--omur", type=int, default=4, help="Satırda yazmıyorsa yenileme ömrü, yıl (varsayılan 4)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "donanim_zimmet_envanteri.xlsx")
    a = p.parse_args(argv)
    ornek = a.envanter is None
    yol, personel = (ORNEK / "envanter.csv", a.personel or ORNEK / "personel.csv") if ornek else (a.envanter, a.personel)
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    if bugun is None or not 0 <= a.gun <= 365 or not 1 <= a.omur <= 20:
        print("[X] --tarih GG.AA.YYYY biçiminde, --gun 0–365, --omur 1–20 arasında olmalı")
        return 2
    for y in (yol, personel):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(yol, a.cikti, bugun, a.gun, a.omur, personel)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    say = lambda onem: sum(u["onem"] == onem for u in r["uyarilar"])   # noqa: E731
    print(f"[OK] {len(r['cihazlar'])} cihaz · {len(r['kisiler'])} zimmetli kişi · uyarı: {say('Yüksek')} yüksek, {say('Orta')} orta, {say('Bilgi')} bilgi")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
