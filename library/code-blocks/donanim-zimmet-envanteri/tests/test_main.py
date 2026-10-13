import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)


def c(no, tur="Dizüstü", kisi="Ali Can", durum="Kullanımda", alim=None, garanti=None, seri="", omur=None, departman="BT"):
    return {"no": no, "tur": tur, "model": "M", "seri": seri, "kisi": kisi, "departman": departman, "zimmet": None, "alim": alim, "garanti": garanti, "durum": durum, "omur": omur}


class DenetimTesti(unittest.TestCase):
    def test_garanti_ve_yenileme(self):
        l = [c("1", garanti=date(2026, 10, 11)), c("2", kisi="B", garanti=date(2026, 10, 12)), c("3", kisi="C", garanti=date(2026, 12, 11)), c("4", kisi="D", garanti=date(2026, 12, 12)), c("5", kisi="E"),
             c("6", kisi="F", durum="Hurda", garanti=date(2020, 1, 1), alim=date(2015, 1, 1)), c("7", kisi="", durum="Arızalı", garanti=date(2027, 1, 1))]
        r = main.denetle(l, BUGUN)
        self.assertEqual([x["garanti_durum"] for x in r["cihazlar"]], ["Bitmiş", "Bitiyor", "Bitiyor", "Devam ediyor", "", "", "Devam ediyor"])
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Arızalı cihaz garanti kapsamında", "7"), ("Garantisi bitiyor", "2"), ("Garantisi bitiyor", "3"), ("Garantisi bitmiş", "1"), ("Zimmet düşülmemiş", "6")])
        l = [c("1", alim=date(2022, 10, 12)), c("2", kisi="B", alim=date(2022, 10, 13)), c("3", kisi="C", alim=date(2022, 12, 11)), c("4", kisi="D", alim=date(2022, 12, 12)), c("5", kisi="E", alim=date(2024, 2, 29), omur=3),
             c("6", kisi="F", alim=date(2025, 10, 12), omur=1), c("7", kisi="G", alim=date(2022, 10, 12), omur=6)]
        r = main.denetle(l, BUGUN)
        self.assertEqual([(x["yenileme"], x["yenileme_durum"], x["yas"]) for x in r["cihazlar"]],
                         [(date(2026, 10, 12), "Zamanı gelmiş", D("4.0")), (date(2026, 10, 13), "Yaklaşıyor", D("4.0")), (date(2026, 12, 11), "Yaklaşıyor", D("3.8")), (date(2026, 12, 12), "Var", D("3.8")),
                          (date(2027, 2, 28), "Var", D("2.6")), (date(2026, 10, 12), "Zamanı gelmiş", D("1.0")), (date(2028, 10, 12), "Var", D("4.0"))])
        self.assertEqual(main.denetle([c("1", alim=date(2022, 10, 12))], BUGUN, 60, 5)["cihazlar"][0]["yenileme_durum"], "Var")

    def test_zimmet_denetimleri(self):
        l = [c("1", seri="S1"), c("2", seri="s1", kisi="Ece Su"), c("2", tur="Telefon", kisi="Ece Su"), c("4", kisi=""), c("5", kisi="Ali Can", durum="Depoda"), c("6", kisi="ALİ CAN"), c("7", tur="Telefon", kisi="Giden Kişi"),
             c("8", kisi="Giden Kişi", durum="Hurda"), c("9", kisi="", durum="Depoda")]
        r = main.denetle(l, BUGUN, personel={"ali can", "ece su"})
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Bilgi", "Bir kişide aynı türden birden çok cihaz", "Ali Can · Dizüstü"), ("Orta", "Zimmet düşülmemiş", "5"), ("Orta", "Zimmet düşülmemiş", "8"), ("Yüksek", "Mükerrer demirbaş numarası", "2"),
                          ("Yüksek", "Mükerrer seri numarası", "S1"), ("Yüksek", "Zimmetli kişi personel listesinde yok", "7"), ("Yüksek", "Zimmetsiz cihaz", "4")])
        self.assertEqual([(k["kisi"], k["adet"], k["cihazlar"]) for k in r["kisiler"]],
                         [("Ali Can", 3, ["Dizüstü (1)", "Dizüstü (5)", "Dizüstü (6)"]), ("Ece Su", 2, ["Dizüstü (2)", "Telefon (2)"]), ("Giden Kişi", 1, ["Telefon (7)"])])
        self.assertEqual([(t["tur"], t["adet"], t["Kullanımda"], t["Depoda"], t["Hurda"]) for t in r["turler"]], [("Dizüstü", 7, 4, 2, 1), ("Telefon", 2, 2, 0, 0)])
        self.assertNotIn("Zimmetli kişi personel listesinde yok", [u["tur"] for u in main.denetle(l, BUGUN)["uyarilar"]])
        with self.assertRaises(ValueError):
            main.denetle([], BUGUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "e.xlsx"
            r = main.calistir(ORNEK / "envanter.csv", cikti, BUGUN, personel_yol=ORNEK / "personel.csv")
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Yüksek", "Mükerrer seri numarası", "SN-A1002"), ("Yüksek", "Zimmetli kişi personel listesinde yok", "BT-0107"), ("Yüksek", "Zimmetsiz cihaz", "BT-0104"), ("Orta", "Garantisi bitiyor", "BT-0103"),
                              ("Orta", "Garantisi bitiyor", "BT-0104"), ("Orta", "Yenileme zamanı gelmiş", "BT-0102"), ("Orta", "Yenileme zamanı gelmiş", "BT-0109"), ("Orta", "Zimmet düşülmemiş", "BT-0111"),
                              ("Bilgi", "Bir kişide aynı türden birden çok cihaz", "Ayşe Kara · Dizüstü Bilgisayar"), ("Bilgi", "Garantisi bitmiş", "BT-0102"), ("Bilgi", "Garantisi bitmiş", "BT-0105"),
                              ("Bilgi", "Garantisi bitmiş", "BT-0109"), ("Bilgi", "Garantisi bitmiş", "BT-0111")])
            self.assertEqual([(t["tur"], t["adet"], t["garanti_bitmis"], t["yenilenecek"]) for t in r["turler"]],
                             [("Dizüstü Bilgisayar", 6, 1, 1), ("Cep Telefonu", 2, 0, 0), ("Monitör", 2, 2, 0), ("Yazıcı", 1, 1, 1)])
            self.assertEqual([(k["kisi"], k["adet"]) for k in r["kisiler"]], [("Ayşe Kara", 3), ("Burak Şen", 2), ("Deniz Acar", 2), ("Emre Tunç", 1), ("Fatma Öz", 1)])
            self.assertEqual((r["cihazlar"][1]["yas"], r["cihazlar"][1]["yenileme"]), (D("4.4"), date(2026, 6, 1)))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Envanter", "Kişi Özeti", "Tür Özeti", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, pr = Path(tmp) / "e.csv", Path(tmp) / "p.csv"
            y.write_text("Envanter No;Cihaz Türü;Model;IMEI;Kullanıcı;Birim;Satın Alma Tarihi;Garanti;Durum;Ömür\nE1;Telefon;T;111;Ali Can;Satış;01.02.2025;01.02.2027;;3\nE2;Telefon;T;;;;x;;;\nE3;Tablet;T;;Ece Su;;;;SERVİSTE;\n"
                         "E4;Tablet;T;;;;;;ödünç;\n;Tablet;;;;;;;;\n", encoding="utf-8")
            pr.write_text("Adı Soyadı\nALİ CAN\n\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["tur"], x["seri"], x["kisi"], x["alim"], x["garanti"], x["durum"], x["omur"]) for x in l], [u["tur"] for u in uy]),
                             ([("E1", "Telefon", "111", "Ali Can", date(2025, 2, 1), date(2027, 2, 1), "Kullanımda", 3), ("E2", "Telefon", "", "", None, None, "Depoda", None),
                               ("E3", "Tablet", "", "Ece Su", None, None, "Arızalı", None), ("E4", "Tablet", "", "", None, None, "Depoda", None)], ["Tarih okunamadı", "Durum tanınmadı", "Demirbaş numarası yok"]))
            self.assertEqual(main.personel_oku(pr), {"ali can"})
            c_ = str(Path(tmp) / "x.xlsx")
            arg = ["--envanter", str(y), "--cikti", c_]
            self.assertEqual(main.main(arg + ["--personel", str(pr), "--tarih", "12.10.2026", "--gun", "90", "--omur", "5"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual((main.main(arg + ["--gun", "-1"]), main.main(arg + ["--omur", "0"]), main.main(arg + ["--tarih", "x"])), (2, 2, 2))
            self.assertEqual((main.main(["--envanter", str(Path(tmp) / "yok.csv")]), main.main(arg + ["--personel", str(Path(tmp) / "yok.csv")])), (1, 1))
            for icerik in ("Tür;Durum\nA;\n", "Demirbaş No;Tür\n;A\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c_]), 0)


if __name__ == "__main__":
    unittest.main()
