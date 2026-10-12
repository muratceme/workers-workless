"""
Randevu Doluluk ve Gelmeme Analizi — Workers / Workless kod bloğu
Hasta Hizmetleri › Randevu Planlama Sorumlusu

Hekim ve bölüm bazında randevu doluluk, gelmeme (no-show) ve iptal oranlarını hesaplayıp fazla randevu (overbooking) önerisi çıkarır.
  Gelmeme oranı %      = gelmeyen ÷ (gelen + gelmeyen) × 100              (iptaller paydaya girmez)
  İptal oranı %        = iptal ÷ (gelen + gelmeyen + iptal) × 100
  Kapasite             = günlük randevu kapasitesi × çalışılan gün          (kapasite dosyası verilirse; çalışılan gün yazılmadıysa randevusu olan gün sayısı)
  Doluluk %            = (gelen + gelmeyen) ÷ kapasite × 100;   gerçekleşen doluluk % = gelen ÷ kapasite × 100
  Fazla randevu önerisi (hekim, günlük) = günlük kapasite × p ÷ (1 − p) × emniyet katsayısı, aşağı yuvarlanır;  p = gelmeme oranı
     p ÷ (1 − p) gelmeyenleri ortalamada tam telafi eden orandır; emniyet katsayısı (--emniyet, 0,5) ve üst sınır (--ust-sinir, kapasitenin %10'u) öneriyi temkinli tutar.
     Öneri yalnız (gelen + gelmeyen) ≥ --min-randevu (30), gelmeme ≥ --min-gelmeme (%5) ve doluluk ≥ --min-doluluk (%90) ise verilir.
Öneri istatistiksel bir ortalamadır: bazı günler herkes gelir ve bekleme uzar. Karar başhekimlik ve hekimle birlikte verilmelidir. Hasta kimlik bilgisi gerekmez.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"tarih": ("tarih", "randevu tarihi", "gun"), "bolum": ("bolum", "brans", "poliklinik", "klinik"), "hekim": ("hekim", "doktor", "hekim adi"), "durum": ("durum", "randevu durumu", "sonuc"),
         "adet": ("adet", "randevu sayisi", "sayi")}
KAPASITE = {"hekim": SUTUN["hekim"], "gunluk": ("gunluk kapasite", "gunluk randevu kapasitesi", "kapasite", "gunluk slot"), "gun": ("calisilan gun", "gun sayisi", "calisma gunu")}
GUNLER = ("Pazartesi", "Salı", "Çarşamba", "Perşembe", "Cuma", "Cumartesi", "Pazar")
Z = Decimal(0)


def durum_coz(x) -> str | None:
    k = katla(x)
    if "iptal" in k:
        return "iptal"
    if "gelmedi" in k or "no show" in k or "noshow" in k or "gelmeyen" in k:
        return "gelmedi"
    if any(a in k for a in ("geldi", "muayene", "tamamlandi", "gelen")):
        return "geldi"
    return None


def oku(yol: Path, kapasite_yol: Path | None = None) -> tuple[list[dict], dict, list[dict]]:
    uy, l, kap, taninmayan = [], [], {}, defaultdict(int)
    for r in kayitlar(yol, SUTUN, ("tarih", "hekim", "durum")):
        h = metin(r["hekim"])
        if not h:
            continue
        kim = f"{h} (satır {r['_satir']})"
        t, d, adet = tarih(r["tarih"]), durum_coz(r["durum"]), para(r["adet"]) if metin(r["adet"]) else Decimal(1)
        if t is None or adet is None or adet < 0:
            uy.append(uyari("Yüksek", "Tarih veya adet okunamadı", kim, f"Tarih '{metin(r['tarih'])}', adet '{metin(r['adet'])}'. Satır alınmadı."))
            continue
        if d is None:
            taninmayan[metin(r["durum"]) or "(boş)"] += 1
            continue
        l.append({"tarih": t, "bolum": metin(r["bolum"]) or "Belirtilmemiş", "hekim": h, "durum": d, "adet": adet})
    for d, n in taninmayan.items():
        uy.append(uyari("Orta", "Durum tanınmadı", d, f"{n} satır alınmadı. Tanınan durumlar: Geldi, Gelmedi, İptal (ileri tarihli veya bekleyen randevular analize girmez)."))
    if kapasite_yol:
        for r in kayitlar(kapasite_yol, KAPASITE, ("hekim", "gunluk")):
            h, g, gun = metin(r["hekim"]), para(r["gunluk"]), para(r["gun"])
            if not h:
                continue
            if g is None or g <= 0 or (gun is not None and gun <= 0):
                uy.append(uyari("Yüksek", "Kapasite okunamadı", h, f"{kapasite_yol.name} satır {r['_satir']}: günlük kapasite '{metin(r['gunluk'])}', çalışılan gün '{metin(r['gun'])}'. Bu hekim için doluluk hesaplanmadı."))
                continue
            kap[katla(h)] = {"gunluk": g, "gun": gun}
    return l, kap, uy


def analiz(randevular: list[dict], kapasite: dict | None = None, min_randevu: Decimal = Decimal(30), min_gelmeme: Decimal = Decimal(5), min_doluluk: Decimal = Decimal(90),
           emniyet: Decimal = Decimal("0.5"), ust_sinir: Decimal = Decimal(10)) -> dict:
    if not randevular:
        raise ValueError("Analiz edilecek randevu yok.")
    kapasite = kapasite or {}

    def oranla(x: dict) -> dict:
        gerceklesen, toplam = x["geldi"] + x["gelmedi"], x["geldi"] + x["gelmedi"] + x["iptal"]
        k = x.get("kapasite")
        return x | {"randevu": toplam, "gelmeme": yuvarla(x["gelmedi"] / gerceklesen * 100) if gerceklesen else None, "iptal_orani": yuvarla(x["iptal"] / toplam * 100) if toplam else None,
                    "doluluk": yuvarla(gerceklesen / k * 100) if k else None, "gercek_doluluk": yuvarla(x["geldi"] / k * 100) if k else None}

    def grupla(anahtar, ad) -> dict:
        g = {}
        for r in randevular:
            x = g.setdefault(anahtar(r), ad(r) | {"geldi": Z, "gelmedi": Z, "iptal": Z, "gunler": set()})
            x[r["durum"]] += r["adet"]
            x["gunler"].add(r["tarih"])
        return g

    hekimler = grupla(lambda r: katla(r["hekim"]), lambda r: {"hekim": r["hekim"], "bolum": r["bolum"]})
    for a, h in hekimler.items():
        k = kapasite.get(a)
        h["gunluk"] = k["gunluk"] if k else None
        h["calisilan"] = (k["gun"] or len(h["gunler"])) if k else len(h["gunler"])
        h["kapasite"] = h["gunluk"] * h["calisilan"] if k else None
        h |= oranla(h)
        h["oneri"], gerceklesen = 0, h["geldi"] + h["gelmedi"]
        if not k:
            h["gerekce"] = "Kapasite verilmedi"
        elif gerceklesen < min_randevu:
            h["gerekce"] = f"Veri az ({gerceklesen:f} randevu)"
        elif h["gelmeme"] < min_gelmeme:
            h["gerekce"] = "Gelmeme oranı düşük"
        elif h["doluluk"] < min_doluluk:
            h["gerekce"] = "Doluluk düşük: önce boş randevular doldurulmalı"
        else:
            p = h["gelmedi"] / gerceklesen
            tam = h["gunluk"] * p / (1 - p) if p < 1 else Z
            h["oneri"] = int(min(math.floor(tam * emniyet), math.floor(h["gunluk"] * ust_sinir / 100)))
            h["gerekce"] = f"Tam telafi {yuvarla(tam)} randevu/gün; emniyet katsayısı ve üst sınır uygulandı" if h["oneri"] else "Hesaplanan ek randevu 1'in altında"
    bolumler = grupla(lambda r: katla(r["bolum"]), lambda r: {"bolum": r["bolum"]})
    for a, b in bolumler.items():
        hk = [h for h in hekimler.values() if katla(h["bolum"]) == a]
        b["kapasite"] = sum((h["kapasite"] for h in hk), Z) if hk and all(h["kapasite"] for h in hk) else None
        b["hekim"] = len(hk)
        b |= oranla(b)
    gunler = grupla(lambda r: r["tarih"].weekday(), lambda r: {"gun": GUNLER[r["tarih"].weekday()], "no": r["tarih"].weekday()})
    genel = {"geldi": sum((h["geldi"] for h in hekimler.values()), Z), "gelmedi": sum((h["gelmedi"] for h in hekimler.values()), Z), "iptal": sum((h["iptal"] for h in hekimler.values()), Z),
             "kapasite": sum((h["kapasite"] for h in hekimler.values()), Z) if all(h["kapasite"] for h in hekimler.values()) else None}
    return {"hekimler": sorted(hekimler.values(), key=lambda h: (-(h["gelmeme"] or Z), h["hekim"])), "bolumler": sorted(bolumler.values(), key=lambda b: (-(b["gelmeme"] or Z), b["bolum"])),
            "gunler": [oranla(g) for g in sorted(gunler.values(), key=lambda g: g["no"])], "genel": oranla(genel), "ayar": (min_randevu, min_gelmeme, min_doluluk, emniyet, ust_sinir),
            "kapasitesiz": sorted(k for k in kapasite if k not in hekimler)}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    mr, mg, md, em, us = r["ayar"]
    bas = ["Randevu", "Gelen", "Gelmeyen", "İptal", "Gelmeme %", "İptal %", "Kapasite", "Doluluk %", "Gerçekleşen Doluluk %"]
    deger = lambda x: [x["randevu"], x["geldi"], x["gelmedi"], x["iptal"], x["gelmeme"], x["iptal_orani"], x.get("kapasite"), x.get("doluluk"), x.get("gercek_doluluk")]   # noqa: E731
    ws = wb.active
    ws.title = "Hekim Bazında"
    _baslik(ws, ["Hekim", "Bölüm", *bas, "Günlük Kapasite", "Çalışılan Gün", "Önerilen Ek Randevu / Gün", "Gerekçe", "Karar (siz doldurun)"], (26, 20, 9, 8, 9, 8, 10, 8, 9, 10, 12, 10, 10, 12, 52, 24))
    for h in r["hekimler"]:
        _satir(ws, [h["hekim"], h["bolum"], *deger(h), h["gunluk"], h["calisilan"], h["oneri"] or None, h["gerekce"], ""], para_s=(7, 8, 10, 11), dolgu={14: YESIL if h["oneri"] else None}, sar=(15,))
    g = r["genel"]
    _satir(ws, ["Toplam", "", *deger(g)], para_s=(7, 8, 10, 11))
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.append([])
    ws.append([f"Öneri = günlük kapasite × p ÷ (1 − p) × {em.normalize():f}, en çok kapasitenin %{us.normalize():f}'i; p = gelmeme oranı. Koşullar: en az {mr.normalize():f} randevu, gelmeme ≥ %{mg.normalize():f}, "
               f"doluluk ≥ %{md.normalize():f}. Öneri ortalamaya dayanır; herkesin geldiği günlerde bekleme uzar. Karar hekim ve yönetimle birlikte verilmelidir."])
    ws = wb.create_sheet("Bölüm Bazında")
    _baslik(ws, ["Bölüm", "Hekim Sayısı", *bas], (24, 9, 9, 8, 9, 8, 10, 8, 9, 10, 12))
    for b in r["bolumler"]:
        _satir(ws, [b["bolum"], b["hekim"], *deger(b)], para_s=(7, 8, 10, 11))
    ws = wb.create_sheet("Haftanın Günü")
    _baslik(ws, ["Gün", *bas[:6]], (14, 9, 8, 9, 8, 10, 8))
    en = max((x["gelmeme"] or Z for x in r["gunler"]), default=Z)
    for x in r["gunler"]:
        _satir(ws, [x["gun"], *deger(x)[:6]], para_s=(6, 7), dolgu={6: SARI if x["gelmeme"] and x["gelmeme"] == en and len(r["gunler"]) > 1 else None})
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, kapasite_yol: Path | None = None, **ayar) -> dict:
    randevular, kapasite, uy = oku(yol, kapasite_yol)
    r = analiz(randevular, kapasite, **ayar)
    for k in r["kapasitesiz"]:
        uy.append(uyari("Orta", "Kapasitesi olan hekimin randevusu yok", k, "Kapasite dosyasında var, randevu dosyasında bu adla kayıt yok. Hekim adlarının iki dosyada aynı yazıldığını kontrol edin."))
    if kapasite_yol:
        for h in r["hekimler"]:
            if h["kapasite"] is None:
                uy.append(uyari("Orta", "Hekimin kapasitesi yok", h["hekim"], "Kapasite dosyasında bulunamadı; doluluk ve fazla randevu önerisi hesaplanmadı."))
            elif h["doluluk"] > 100:
                uy.append(uyari("Orta", "Doluluk %100'ün üzerinde", h["hekim"], f"Doluluk %{h['doluluk']}. Günlük kapasite veya çalışılan gün eksik girilmiş olabilir; ek randevu zaten veriliyor olabilir."))
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Randevu doluluk, gelmeme ve iptal oranlarını hesaplar; fazla randevu önerir.")
    p.add_argument("--randevular", type=Path, default=ORNEK / "randevular.csv", help="Tarih, Bölüm, Hekim, Durum (Geldi / Gelmedi / İptal), Adet (isteğe bağlı)")
    p.add_argument("--kapasite", type=Path, help="İsteğe bağlı: Hekim, Günlük Kapasite, Çalışılan Gün")
    p.add_argument("--min-randevu", default="30", help="Öneri için en az gerçekleşen randevu sayısı (varsayılan 30)")
    p.add_argument("--min-gelmeme", default="5", help="Öneri için en az gelmeme oranı %% (varsayılan 5)")
    p.add_argument("--min-doluluk", default="90", help="Öneri için en az doluluk %% (varsayılan 90)")
    p.add_argument("--emniyet", default="0,5", help="Tam telafi oranına uygulanan emniyet katsayısı, 0-1 (varsayılan 0,5)")
    p.add_argument("--ust-sinir", default="10", help="Ek randevu günlük kapasitenin en çok yüzde kaçı olabilir (varsayılan 10)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "randevu_analizi.xlsx")
    a = p.parse_args(argv)
    kap = a.kapasite or (ORNEK / "kapasite.csv" if a.randevular == ORNEK / "randevular.csv" else None)
    d = [para(x) for x in (a.min_randevu, a.min_gelmeme, a.min_doluluk, a.emniyet, a.ust_sinir)]
    if any(x is None or x < 0 for x in d) or not 0 < d[3] <= 1:
        print("[X] Eşikler sıfır veya pozitif sayı; --emniyet 0 ile 1 arasında (0 hariç) olmalı")
        return 2
    for y in (a.randevular, kap):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.randevular, a.cikti, kap, min_randevu=d[0], min_gelmeme=d[1], min_doluluk=d[2], emniyet=d[3], ust_sinir=d[4])
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    g = r["genel"]
    print(f"[OK] {g['randevu']:f} randevu · gelmeme %{g['gelmeme']} · iptal %{g['iptal_orani']}" + (f" · doluluk %{g['doluluk']} (gerçekleşen %{g['gercek_doluluk']})" if g["doluluk"] is not None else ""))
    for h in r["hekimler"]:
        if h["oneri"]:
            print(f"[!] {h['hekim']} ({h['bolum']}): gelmeme %{h['gelmeme']}, doluluk %{h['doluluk']} → günde {h['oneri']} ek randevu önerisi")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
