import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
PZT = date(2026, 10, 5)


def rv(geldi, gelmedi, iptal=0, hekim="H", bolum="B", gun=PZT):
    return [{"tarih": gun, "bolum": bolum, "hekim": hekim, "durum": d, "adet": D(n)} for d, n in (("geldi", geldi), ("gelmedi", gelmedi), ("iptal", iptal)) if n]


def kap(gunluk, gun=None, hekim="h"):
    return {hekim: {"gunluk": D(gunluk), "gun": None if gun is None else D(gun)}}


def tek(geldi, gelmedi, iptal=0, kapasite=None, **kw):
    return main.analiz(rv(geldi, gelmedi, iptal), kapasite, **kw)["hekimler"][0]


class AnalizTesti(unittest.TestCase):
    def test_oranlar(self):
        h = tek(75, 20, 10, kap(20, 5))
        self.assertEqual((h["randevu"], h["gelmeme"], h["iptal_orani"], h["kapasite"], h["doluluk"], h["gercek_doluluk"]), (D(105), D("21.05"), D("9.52"), D(100), D("95.00"), D("75.00")))
        h = tek(0, 0, 5)
        self.assertEqual((h["gelmeme"], h["iptal_orani"], h["doluluk"], h["gerekce"]), (None, D("100.00"), None, "Kapasite verilmedi"))
        self.assertEqual(tek(10, 2, kapasite=kap(12))["kapasite"], D(12))                                # çalışılan gün yazılmadı: randevusu olan gün sayısı (1)

    def test_fazla_randevu_onerisi(self):
        self.assertEqual((tek(75, 20, 10, kap(20, 5))["oneri"], tek(50, 10, 0, kap(12, 5))["oneri"]), (2, 1))       # 20 × 0,2105 ÷ 0,7895 × 0,5 = 2,67 → 2;  12 × 0,2 × 0,5 = 1,2 → 1
        self.assertEqual(tek(60, 40, 0, kap(20, 5))["oneri"], 2)                                                    # 20 × 0,4 ÷ 0,6 × 0,5 = 6,67 ama üst sınır %10 = 2
        self.assertEqual(tek(60, 40, 0, kap(20, 5), ust_sinir=D(50), emniyet=D(1))["oneri"], 10)                    # üst sınır 10, tam telafi 13,33
        self.assertEqual([tek(g, m, 0, kap(20, 5))["gerekce"].split(":")[0].split(" (")[0] for g, m in ((20, 9), (96, 4), (70, 19), (72, 18))],
                         ["Veri az", "Gelmeme oranı düşük", "Doluluk düşük", "Tam telafi 5.00 randevu/gün; emniyet katsayısı ve üst sınır uygulandı"])
        self.assertEqual([tek(g, 100 - g, 0, kap(20, 5))["oneri"] for g in (95, 94)], [0, 0])                       # tam %5 koşulu sağlar ama 20 × 0,0526 × 0,5 < 1
        self.assertEqual(tek(95, 5, 0, kap(20, 5))["gerekce"], "Hesaplanan ek randevu 1'in altında")
        self.assertEqual(tek(81, 9, 0, kap(20, 5))["gerekce"], "Tam telafi 2.22 randevu/gün; emniyet katsayısı ve üst sınır uygulandı")    # doluluk tam %90
        self.assertEqual(tek(25, 5, 0, kap(30, 1), min_randevu=D(10))["oneri"], 3)

    def test_bolum_gun_ve_genel(self):
        l = rv(11, 8, 2, "A", "Dahiliye") + rv(16, 3, 2, "A", "Dahiliye", date(2026, 10, 6)) + rv(11, 1, 1, "B", "Dahiliye") + rv(10, 2, 0, "C", "Kardiyoloji")
        r = main.analiz(l, kap(20, hekim="a") | kap(16, 1, "b") | kap(99, hekim="x"))
        self.assertEqual([(b["bolum"], b["hekim"], b["randevu"], b["gelmeme"], b["kapasite"], b["doluluk"]) for b in r["bolumler"]], [("Dahiliye", 2, D(55), D("24.00"), D(56), D("89.29")), ("Kardiyoloji", 1, D(12), D("16.67"), None, None)])
        self.assertEqual([(g["gun"], g["geldi"], g["gelmedi"], g["gelmeme"]) for g in r["gunler"]], [("Pazartesi", D(32), D(11), D("25.58")), ("Salı", D(16), D(3), D("15.79"))])
        self.assertEqual((r["genel"]["randevu"], r["genel"]["gelmeme"], r["genel"]["iptal_orani"], r["genel"]["kapasite"], r["kapasitesiz"]), (D(67), D("22.58"), D("7.46"), None, ["x"]))
        self.assertEqual([h["hekim"] for h in r["hekimler"]], ["A", "C", "B"])
        with self.assertRaises(ValueError):
            main.analiz([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "r.xlsx"
            r = main.calistir(ORNEK / "randevular.csv", cikti, ORNEK / "kapasite.csv")
            self.assertEqual([(h["hekim"], h["gelmeme"], h["iptal_orani"], h["doluluk"], h["gercek_doluluk"], h["oneri"]) for h in r["hekimler"]],
                             [("Dr. A", D("21.05"), D("9.52"), D("95.00"), D("75.00"), 2), ("Dr. C", D("16.67"), D("0.00"), D("100.00"), D("83.33"), 1), ("Dr. B", D("8.33"), D("7.69"), D("75.00"), D("68.75"), 0)])
            g = r["genel"]
            self.assertEqual((g["geldi"], g["gelmedi"], g["iptal"], g["gelmeme"], g["iptal_orani"], g["kapasite"], g["doluluk"], g["gercek_doluluk"]),
                             (D(180), D(35), D(15), D("16.28"), D("6.52"), D(240), D("89.58"), D("75.00")))
            self.assertEqual([(b["bolum"], b["gelmeme"], b["doluluk"]) for b in r["bolumler"]], [("Kardiyoloji", D("16.67"), D("100.00")), ("Dahiliye", D("16.13"), D("86.11"))])
            self.assertEqual([(x["gun"], x["gelmeme"]) for x in r["gunler"]], [("Pazartesi", D("25.58")), ("Salı", D("13.95")), ("Çarşamba", D("13.95")), ("Perşembe", D("13.95")), ("Cuma", D("13.95"))])
            self.assertEqual((r["uyarilar"], load_workbook(cikti).sheetnames), ([], ["Hekim Bazında", "Bölüm Bazında", "Haftanın Günü", "Uyarılar"]))

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, k = Path(tmp) / "r.csv", Path(tmp) / "k.csv"
            y.write_text("Randevu Tarihi;Doktor;Durum\n05.10.2026;A;Muayene edildi\n05.10.2026;A;GELMEDİ\n05.10.2026;A;Hasta iptal etti\n05.10.2026;A;Bekliyor\n05.10.2026;A;\nbugün;A;Geldi\n05.10.2026;;Geldi\n",
                         encoding="utf-8")
            k.write_text("Hekim;Günlük Kapasite;Çalışılan Gün\nA;2;1\nB;çok;\nC;10;0\n", encoding="utf-8")
            l, kp, uy = main.oku(y, k)
            self.assertEqual(([(x["durum"], x["adet"], x["bolum"]) for x in l], kp), ([("geldi", D(1), "Belirtilmemiş"), ("gelmedi", D(1), "Belirtilmemiş"), ("iptal", D(1), "Belirtilmemiş")], {"a": {"gunluk": D(2), "gun": D(1)}}))
            self.assertEqual(sorted(u["tur"] for u in uy), ["Durum tanınmadı", "Durum tanınmadı", "Kapasite okunamadı", "Kapasite okunamadı", "Tarih veya adet okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--randevular", str(y), "--cikti", c]
            k.write_text("Hekim;Günlük Kapasite\nA;1\nZ;5\n", encoding="utf-8")
            r = main.calistir(y, Path(c), k)
            self.assertEqual([u["tur"] for u in r["uyarilar"]], ["Tarih veya adet okunamadı", "Doluluk %100'ün üzerinde","Durum tanınmadı", "Durum tanınmadı", "Kapasitesi olan hekimin randevusu yok"])
            self.assertEqual(main.main(arg + ["--kapasite", str(k), "--min-randevu", "1", "--emniyet", "1", "--ust-sinir", "50"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(arg + ["--emniyet", "1,5"]), 2)
            self.assertEqual(main.main(arg + ["--kapasite", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Tarih;Hekim;Durum\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
