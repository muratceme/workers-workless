import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)


def g(n):
    return None if n is None else BUGUN + timedelta(days=n)


def b(no, termin=30, yas=10, kapanis=None, onem="Orta", birim="B", revizyon=0, tekrar=False):
    """termin ve kapanis: bugünden kaç gün sonra (eksi = geçmiş); yas: açılış kaç gün önce."""
    return {"no": no, "ad": "Bulgu", "birim": birim, "onem": onem, "acilis": g(-yas), "termin": g(termin), "kapanis": g(kapanis), "sahip": "S", "revizyon": revizyon, "tekrar": tekrar}


class DegerlendirTesti(unittest.TestCase):
    def test_durumlar(self):
        l = [b("dün", -1), b("bugün", 0), b("onbeş", 15), b("onaltı", 16), b("terminsiz", None), b("zamanında", -5, 40, -5), b("geç", -5, 40, -4), b("terminsiz_kapalı", None, 40, -1), b("yüksek", -20, 100, onem="Yüksek"),
             b("ertelenen", 40, revizyon=2), b("tekrar", 40, tekrar=True), b("kapalı_tekrar", -5, 40, -9, revizyon=5, tekrar=True)]
        r = main.degerlendir(l, BUGUN)
        self.assertEqual({x["no"]: (x["durum"], x["gecikme"], x["yas"], x["sure"]) for x in r["bulgular"]},
                         {"dün": ("Gecikmiş", 1, 10, None), "bugün": ("Yaklaşıyor", None, 10, None), "onbeş": ("Yaklaşıyor", None, 10, None), "onaltı": ("Açık", None, 10, None), "terminsiz": ("Termin yok", None, 10, None),
                          "zamanında": ("Zamanında kapandı", None, None, 35), "geç": ("Geç kapandı", 1, None, 36), "terminsiz_kapalı": ("Zamanında kapandı", None, None, 39), "yüksek": ("Gecikmiş", 20, 100, None),
                          "ertelenen": ("Açık", None, 10, None), "tekrar": ("Açık", None, 10, None), "kapalı_tekrar": ("Zamanında kapandı", None, None, 31)})
        self.assertEqual([x["no"] for x in r["acik"]], ["yüksek", "dün", "terminsiz", "bugün", "onbeş", "ertelenen", "onaltı", "tekrar"])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         sorted([("Orta", "Gecikmiş bulgu", "dün · B"), ("Bilgi", "Termin yaklaşıyor", "bugün · B"), ("Bilgi", "Termin yaklaşıyor", "onbeş · B"), ("Orta", "Termin yok", "terminsiz · B"),
                                 ("Yüksek", "Gecikmiş yüksek önemli bulgu", "yüksek · B"), ("Orta", "Termin birden çok kez ertelendi", "ertelenen · B"), ("Orta", "Tekrar eden bulgu", "tekrar · B")]))
        self.assertEqual(r["ozet"], {"bulgu": 12, "acik": 8, "gecikmis": 2, "yuksek_acik": 1, "kapali": 4, "zamaninda": 3, "kapanma": D("33.33"), "zamaninda_oran": D("75.00"), "ort_sure": D("35.25"), "ort_yas": D("21.25"), "ort_gecikme": D("10.50")})
        d = lambda **kw: main.degerlendir([b("a", 20, revizyon=1)], BUGUN, **kw)   # noqa: E731
        self.assertEqual((d()["bulgular"][0]["durum"], d(gun=20)["bulgular"][0]["durum"], [u["tur"] for u in d(revizyon=1)["uyarilar"]], d()["uyarilar"]), ("Açık", "Yaklaşıyor", ["Termin birden çok kez ertelendi"], []))
        self.assertEqual([u["tur"] for u in main.degerlendir([b("a", 5, 10, -20)], BUGUN)["uyarilar"]], ["Kapanış açılıştan önce"])
        with self.assertRaises(ValueError):
            main.degerlendir([], BUGUN)

    def test_yas_ve_kirilim(self):
        l = [b("1", 50, 30, birim="X", onem="Yüksek"), b("2", 50, 31, birim="x"), b("3", -3, 90, birim="Y", onem="Düşük"), b("4", 50, 91, birim="Y", onem="Yüksek"), b("5", 50, 180, birim="Y", onem="Belirsiz"), b("6", 50, 181, birim="Z"),
             b("7", -10, 60, -20, birim="Z"), b("8", -10, 60, -5, birim="Z", onem="Yüksek")]
        r = main.degerlendir(l, BUGUN)
        self.assertEqual(({x["no"]: x["yas_dilim"] for x in r["acik"]}, r["yaslar"]),
                         ({"1": "0–30", "2": "31–90", "3": "31–90", "4": "91–180", "5": "91–180", "6": "181+"},
                          [{"dilim": "0–30", "bulgu": 1, "yuksek": 1}, {"dilim": "31–90", "bulgu": 2, "yuksek": 0}, {"dilim": "91–180", "bulgu": 2, "yuksek": 1}, {"dilim": "181+", "bulgu": 1, "yuksek": 0}]))
        k = lambda alan: [(s["ad"], s["bulgu"], s["acik"], s["gecikmis"], s["yuksek_acik"], s["kapali"], s["kapanma"], s["zamaninda_oran"], s["ort_sure"], s["ort_yas"]) for s in r[alan]]   # noqa: E731
        self.assertEqual(k("birimler"), [("X", 2, 2, 0, 1, 0, D("0.00"), None, None, D("30.50")), ("Y", 3, 3, 1, 1, 0, D("0.00"), None, None, D("120.33")), ("Z", 3, 1, 0, 0, 2, D("66.67"), D("50.00"), D("47.50"), D("181.00"))])
        self.assertEqual([x[:6] for x in k("onemler")], [("Yüksek", 3, 2, 0, 2, 1), ("Orta", 3, 2, 0, 0, 1), ("Düşük", 1, 1, 1, 0, 0), ("Belirsiz", 1, 1, 0, 0, 0)])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "bulgular.csv", cikti, BUGUN)
            self.assertEqual({x["no"]: (x["durum"], x["gecikme"], x["yas"], x["sure"]) for x in r["bulgular"]},
                             {"B-01": ("Gecikmiş", 104, 216, None), "B-02": ("Açık", None, 150, None), "B-03": ("Zamanında kapandı", None, None, 108), "B-04": ("Geç kapandı", 15, None, 148), "B-05": ("Gecikmiş", 12, 280, None),
                              "B-06": ("Yaklaşıyor", None, 122, None), "B-07": ("Termin yok", None, 103, None), "B-08": ("Gecikmiş", 27, 55, None), "B-09": ("Zamanında kapandı", None, None, 56)})
            self.assertEqual([x["no"] for x in r["acik"]], ["B-01", "B-05", "B-08", "B-07", "B-06", "B-02"])
            self.assertEqual(r["ozet"], {"bulgu": 9, "acik": 6, "gecikmis": 3, "yuksek_acik": 3, "kapali": 3, "zamaninda": 2, "kapanma": D("33.33"), "zamaninda_oran": D("66.67"), "ort_sure": D("104.00"), "ort_yas": D("154.33"), "ort_gecikme": D("47.67")})
            self.assertEqual([(y["dilim"], y["bulgu"], y["yuksek"]) for y in r["yaslar"]], [("0–30", 0, 0), ("31–90", 1, 0), ("91–180", 3, 1), ("181+", 2, 2)])
            self.assertEqual([(s["ad"], s["bulgu"], s["acik"], s["gecikmis"], s["yuksek_acik"], s["kapali"]) for s in r["birimler"]],
                             [("Bilgi Teknolojileri", 2, 2, 1, 2, 0), ("İdari İşler", 2, 1, 0, 0, 1), ("Muhasebe", 2, 1, 0, 0, 1), ("Satın Alma", 1, 0, 0, 0, 1), ("Şube Operasyon", 2, 2, 2, 1, 0)])
            self.assertEqual([(s["ad"], s["bulgu"], s["acik"], s["gecikmis"], s["kapali"]) for s in r["onemler"]], [("Yüksek", 3, 3, 2, 0), ("Orta", 3, 1, 1, 2), ("Düşük", 3, 2, 0, 1)])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Yüksek", "Gecikmiş yüksek önemli bulgu", "B-01 · Şube Operasyon"), ("Yüksek", "Gecikmiş yüksek önemli bulgu", "B-05 · Bilgi Teknolojileri"), ("Orta", "Gecikmiş bulgu", "B-08 · Şube Operasyon"),
                              ("Orta", "Tekrar eden bulgu", "B-01 · Şube Operasyon"), ("Orta", "Tekrar eden bulgu", "B-07 · Muhasebe"), ("Orta", "Termin birden çok kez ertelendi", "B-01 · Şube Operasyon"),
                              ("Orta", "Termin birden çok kez ertelendi", "B-08 · Şube Operasyon"), ("Orta", "Termin yok", "B-07 · Muhasebe"), ("Bilgi", "Termin yaklaşıyor", "B-06 · İdari İşler")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Açık Bulgular", "Özet", "Birimler", "Önem Dereceleri", "Tüm Bulgular", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "b.csv"
            y.write_text("No;Tespit;Bölüm;Derece;Tespit Tarihi;Hedef Tarih;Tamamlanma Tarihi;Sorumlu;Erteleme Sayısı;Tekrarlayan\nK1;Metin;;KRİTİK;01.03.2026;x;;Ali;2;EVET\nK2;;B;önemsiz;02.03.2026;;05.03.2026;;;\nk1;;;;03.03.2026;;;;;\n"
                         "K3;;;;x;;;;;\nK4;;;;03.03.2026;;;;1,5;\n;;;;03.03.2026;;;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["ad"], x["birim"], x["onem"], x["acilis"], x["termin"], x["kapanis"], x["sahip"], x["revizyon"], x["tekrar"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("K1", "Metin", "(birim yok)", "Yüksek", date(2026, 3, 1), None, None, "Ali", 2, True), ("K2", "", "B", "Belirsiz", date(2026, 3, 2), None, date(2026, 3, 5), "", 0, False)],
                              [("Tarih okunamadı", "K1"), ("Bulgu okunamadı", "k1"), ("Bulgu okunamadı", "K3"), ("Bulgu okunamadı", "K4")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--bulgular", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "12.10.2026", "--gun", "30", "--revizyon", "1"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--gun", "-1"], ["--revizyon", "0"])], [2] * 3)
            self.assertEqual(main.main(["--bulgular", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Bulgu No;Bulgu\nA;B\n", "Bulgu No;Açılış Tarihi\n;01.03.2026\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
