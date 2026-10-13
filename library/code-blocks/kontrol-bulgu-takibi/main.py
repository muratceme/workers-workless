"""
Kontrol Bulgu Takibi — Workers / Workless kod bloğu
İç Kontrol ve Risk Yönetimi › İç Kontrol Uzmanı

İç kontrol bulgularını sorumlu birim, aksiyon ve kapanış tarihine göre izler; gecikenleri raporlar.
  Kapalı bulgu        = kapanış tarihi dolu;  Zamanında kapandı (kapanış ≤ termin ya da termin yok) · Geç kapandı
  Açık bulgu          = Gecikmiş (termin < durum tarihi) · Termin yok · Yaklaşıyor (termine --gun, 15 gün veya az) · Açık
  Gecikme (gün)       = durum tarihi − termin   (açık bulgu);   kapanış − termin (geç kapanan bulgu)
  Yaş (gün)           = durum tarihi − açılış tarihi   (açık bulgu);   yaş dilimleri 0–30, 31–90, 91–180, 181+
  Kapanma oranı %     = kapalı ÷ toplam × 100;   zamanında kapanma % = zamanında kapanan ÷ kapalı × 100
  Ortalama kapanma süresi = kapalı bulgularda (kapanış − açılış) ortalaması
  Çok ertelenen bulgu = termin revizyon sayısı --revizyon (2) ve üstünde olan açık bulgu
Önem derecesini ve termini siz verirsiniz; kod bulguyu değerlendirmez, aksiyonun yeterliliğine bakmaz.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
SUTUN = {"no": ("bulgu no", "bulgu numarasi", "no"), "ad": ("bulgu", "bulgu ozeti", "aciklama", "tespit"), "birim": ("birim", "sorumlu birim", "bolum"), "onem": ("onem", "onem derecesi", "risk seviyesi", "derece"),
         "acilis": ("acilis tarihi", "tespit tarihi", "rapor tarihi"), "termin": ("termin", "termin tarihi", "hedef tarih"), "kapanis": ("kapanis tarihi", "kapanma tarihi", "tamamlanma tarihi"),
         "sahip": ("aksiyon sahibi", "sorumlu", "sorumlu kisi"), "revizyon": ("termin revizyon sayisi", "revizyon sayisi", "erteleme sayisi"), "tekrar": ("tekrar", "tekrar eden", "tekrarlayan")}
ONEM = {"yuksek": "Yüksek", "kritik": "Yüksek", "orta": "Orta", "dusuk": "Düşük"}
ONEM_SIRA = ("Yüksek", "Orta", "Düşük", "Belirsiz")
ACIK_SIRA = ("Gecikmiş", "Termin yok", "Yaklaşıyor", "Açık")
DURUM_RENK = {"Gecikmiş": KIRMIZI, "Termin yok": SARI, "Yaklaşıyor": SARI, "Açık": None, "Geç kapandı": MAVI, "Zamanında kapandı": YESIL}
YAS = ((30, "0–30"), (90, "31–90"), (180, "91–180"), (None, "181+"))
EVET = {"evet", "e", "var", "x", "true", "1"}


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("no", "acilis")):
        no = metin(r["no"])
        if not no:
            continue
        acilis, rev = tarih(r["acilis"]), para(r["revizyon"]) if metin(r["revizyon"]) else Decimal(0)
        if acilis is None or rev is None or rev < 0 or rev != rev.to_integral_value() or katla(no) in gorulen:
            uy.append(uyari("Yüksek", "Bulgu okunamadı", no, f"Satır {r['_satir']}: " + ("aynı bulgu numarası daha önce geçti" if katla(no) in gorulen else f"açılış tarihi '{metin(r['acilis'])}' ya da revizyon sayısı '{metin(r['revizyon'])}' okunamadı")
                            + ". Satır alınmadı."))
            continue
        gorulen.add(katla(no))
        for alan, ad in (("termin", "Termin"), ("kapanis", "Kapanış tarihi")):
            if metin(r[alan]) and tarih(r[alan]) is None:
                uy.append(uyari("Orta", "Tarih okunamadı", no, f"{ad} '{metin(r[alan])}' okunamadı; boş sayıldı."))
        l.append({"no": no, "ad": metin(r["ad"]), "birim": metin(r["birim"]) or "(birim yok)", "onem": ONEM.get(katla(r["onem"]), "Belirsiz"), "acilis": acilis, "termin": tarih(r["termin"]), "kapanis": tarih(r["kapanis"]),
                  "sahip": metin(r["sahip"]), "revizyon": int(rev), "tekrar": katla(r["tekrar"]) in EVET})
    return l, uy


def _ort(d: list[int]) -> Decimal | None:
    return yuvarla(Decimal(sum(d)) / len(d)) if d else None


def _oran(pay: int, payda: int) -> Decimal | None:
    return yuvarla(Decimal(pay * 100) / payda) if payda else None


def degerlendir(bulgular: list[dict], bugun: date, gun: int = 15, revizyon: int = 2) -> dict:
    if not bulgular:
        raise ValueError("İzlenecek bulgu yok.")
    uy = []
    for b in bulgular:
        kim = f"{b['no']} · {b['birim']}"
        b["acik"] = b["kapanis"] is None
        b["yas"] = b["sure"] = b["gecikme"] = b["yas_dilim"] = None
        if b["acik"]:
            b["yas"] = max((bugun - b["acilis"]).days, 0)
            b["yas_dilim"] = next(e for s, e in YAS if s is None or b["yas"] <= s)
            kalan = (b["termin"] - bugun).days if b["termin"] else None
            b["durum"] = "Termin yok" if kalan is None else "Gecikmiş" if kalan < 0 else "Yaklaşıyor" if kalan <= gun else "Açık"
            if b["durum"] == "Gecikmiş":
                b["gecikme"] = -kalan
                uy.append(uyari("Yüksek" if b["onem"] == "Yüksek" else "Orta", "Gecikmiş yüksek önemli bulgu" if b["onem"] == "Yüksek" else "Gecikmiş bulgu", kim,
                                f"Termin {b['termin']:%d.%m.%Y} ({b['gecikme']} gün geçti); bulgu {b['yas']} gündür açık. Aksiyon sahibi: {b['sahip'] or 'yazılmamış'}. {b['ad']}"))
            elif b["durum"] == "Termin yok":
                uy.append(uyari("Orta", "Termin yok", kim, f"Bulgu {b['yas']} gündür açık ama termin verilmemiş. {b['ad']}"))
            elif b["durum"] == "Yaklaşıyor":
                uy.append(uyari("Bilgi", "Termin yaklaşıyor", kim, f"Termin {b['termin']:%d.%m.%Y} ({kalan} gün kaldı). {b['ad']}"))
            if b["revizyon"] >= revizyon:
                uy.append(uyari("Orta", "Termin birden çok kez ertelendi", kim, f"Termin {b['revizyon']} kez revize edilmiş. {b['ad']}"))
            if b["tekrar"]:
                uy.append(uyari("Orta", "Tekrar eden bulgu", kim, f"Daha önce de tespit edilmiş ve hâlâ açık. {b['ad']}"))
        else:
            b["sure"] = max((b["kapanis"] - b["acilis"]).days, 0)
            gec = b["termin"] is not None and b["kapanis"] > b["termin"]
            b["durum"], b["gecikme"] = ("Geç kapandı", (b["kapanis"] - b["termin"]).days) if gec else ("Zamanında kapandı", None)
            if b["kapanis"] < b["acilis"]:
                uy.append(uyari("Orta", "Kapanış açılıştan önce", kim, f"Açılış {b['acilis']:%d.%m.%Y}, kapanış {b['kapanis']:%d.%m.%Y}. Tarihleri kontrol edin."))

    def ozet(liste: list[dict]) -> dict:
        acik, kapali = [b for b in liste if b["acik"]], [b for b in liste if not b["acik"]]
        return {"bulgu": len(liste), "acik": len(acik), "gecikmis": sum(1 for b in acik if b["durum"] == "Gecikmiş"), "yuksek_acik": sum(1 for b in acik if b["onem"] == "Yüksek"), "kapali": len(kapali),
                "zamaninda": sum(1 for b in kapali if b["durum"] == "Zamanında kapandı"), "kapanma": _oran(len(kapali), len(liste)), "zamaninda_oran": _oran(sum(1 for b in kapali if b["durum"] == "Zamanında kapandı"), len(kapali)),
                "ort_sure": _ort([b["sure"] for b in kapali]), "ort_yas": _ort([b["yas"] for b in acik]), "ort_gecikme": _ort([b["gecikme"] for b in acik if b["durum"] == "Gecikmiş"])}

    def kirilim(alan: str, sira=None) -> list[dict]:
        grup = {}
        for b in bulgular:
            grup.setdefault(katla(b[alan]), []).append(b)
        satirlar = [{"ad": g[0][alan], **ozet(g)} for _, g in sorted(grup.items())]
        return sorted(satirlar, key=lambda s: sira.index(s["ad"])) if sira else satirlar

    acik = sorted((b for b in bulgular if b["acik"]), key=lambda b: (ACIK_SIRA.index(b["durum"]), ONEM_SIRA.index(b["onem"]), -(b["gecikme"] or 0), -b["yas"], katla(b["no"])))
    return {"bulgular": bulgular, "acik": acik, "ozet": ozet(bulgular), "birimler": kirilim("birim"), "onemler": kirilim("onem", ONEM_SIRA),
            "yaslar": [{"dilim": e, "bulgu": sum(1 for b in acik if b["yas_dilim"] == e), "yuksek": sum(1 for b in acik if b["yas_dilim"] == e and b["onem"] == "Yüksek")} for _, e in YAS], "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, bugun: date) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Açık Bulgular"
    _baslik(ws, ["Durum", "Bulgu No", "Bulgu", "Birim", "Önem", "Açılış", "Termin", "Gecikme (gün)", "Yaş (gün)", "Aksiyon Sahibi", "Termin Revizyonu", "Tekrar", "Son Durum (siz doldurun)"], (14, 11, 50, 22, 9, 12, 12, 10, 9, 22, 10, 8, 36))
    for b in r["acik"]:
        _satir(ws, [b["durum"], b["no"], b["ad"], b["birim"], b["onem"], b["acilis"], b["termin"], b["gecikme"], b["yas"], b["sahip"], b["revizyon"] or None, "Evet" if b["tekrar"] else None, None], tarih_s=(6, 7),
               dolgu={1: DURUM_RENK[b["durum"]], 5: KIRMIZI if b["onem"] == "Yüksek" else None}, sar=(3,))
    ws = wb.create_sheet("Özet")
    o = r["ozet"]
    _baslik(ws, ["Gösterge", "Değer"], (40, 14))
    for etiket, deger, para_mi in (("Toplam bulgu", o["bulgu"], False), ("Açık bulgu", o["acik"], False), ("Gecikmiş bulgu", o["gecikmis"], False), ("Yüksek önemli açık bulgu", o["yuksek_acik"], False), ("Kapalı bulgu", o["kapali"], False),
                                   ("Kapanma oranı %", o["kapanma"], True), ("Zamanında kapanma %", o["zamaninda_oran"], True), ("Ortalama kapanma süresi (gün)", o["ort_sure"], True), ("Açık bulguların ortalama yaşı (gün)", o["ort_yas"], True),
                                   ("Gecikmişlerde ortalama gecikme (gün)", o["ort_gecikme"], True)):
        _satir(ws, [etiket, deger], para_s=(2,) if para_mi else ())
    ws.append([])
    _baslik(ws, ["Açık bulguların yaşı (gün)", "Bulgu", "Yüksek Önemli"], dondur="A2")
    for y in r["yaslar"]:
        _satir(ws, [y["dilim"], y["bulgu"], y["yuksek"] or None])
    ws.column_dimensions["C"].width = 14
    ws.append([])
    ws.append([f"Durum tarihi: {bugun:%d.%m.%Y}"])
    for baslik, alan, ad in (("Birimler", "birimler", "Birim"), ("Önem Dereceleri", "onemler", "Önem")):
        ws = wb.create_sheet(baslik)
        _baslik(ws, [ad, "Bulgu", "Açık", "Gecikmiş", "Yüksek Önemli Açık", "Kapalı", "Kapanma %", "Zamanında Kapanma %", "Ort. Kapanma Süresi", "Açıkların Ort. Yaşı"], (26, 8, 8, 10, 12, 8, 11, 12, 12, 12))
        for s in r[alan]:
            _satir(ws, [s["ad"], s["bulgu"], s["acik"], s["gecikmis"] or None, s["yuksek_acik"] or None, s["kapali"], s["kapanma"], s["zamaninda_oran"], s["ort_sure"], s["ort_yas"]], para_s=(7, 8, 9, 10), dolgu={4: KIRMIZI if s["gecikmis"] else None})
    ws = wb.create_sheet("Tüm Bulgular")
    _baslik(ws, ["Bulgu No", "Bulgu", "Birim", "Önem", "Açılış", "Termin", "Kapanış", "Durum", "Kapanma Süresi (gün)", "Gecikme (gün)", "Aksiyon Sahibi"], (11, 50, 22, 9, 12, 12, 12, 18, 11, 10, 22))
    for b in r["bulgular"]:
        _satir(ws, [b["no"], b["ad"], b["birim"], b["onem"], b["acilis"], b["termin"], b["kapanis"], b["durum"], b["sure"], b["gecikme"], b["sahip"]], tarih_s=(5, 6, 7), dolgu={8: DURUM_RENK[b["durum"]]}, sar=(2,))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, gun: int = 15, revizyon: int = 2) -> dict:
    bulgular, uy = oku(yol)
    r = degerlendir(bulgular, bugun, gun, revizyon)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, bugun)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="İç kontrol bulgularını birim, aksiyon sahibi ve termine göre izler; gecikenleri raporlar.")
    p.add_argument("--bulgular", type=Path, help="Bulgu No, Bulgu, Birim, Önem, Açılış Tarihi, Termin, Kapanış Tarihi, Aksiyon Sahibi, Termin Revizyon Sayısı, Tekrar")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--gun", type=int, default=15, help="Termine kaç gün kala 'Yaklaşıyor' sayılsın (varsayılan 15)")
    p.add_argument("--revizyon", type=int, default=2, help="Termini bu kadar ve üstü kez ertelenen açık bulgu uyarılır (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kontrol_bulgu_takibi.xlsx")
    a = p.parse_args(argv)
    ornek = a.bulgular is None
    yol = ORNEK / "bulgular.csv" if ornek else a.bulgular
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    if bugun is None or not 0 <= a.gun <= 365 or a.revizyon < 1:
        print("[X] --tarih GG.AA.YYYY; --gun 0–365; --revizyon en az 1 olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, bugun, a.gun, a.revizyon)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {o['bulgu']} bulgu · açık {o['acik']} (gecikmiş {o['gecikmis']}, yüksek önemli {o['yuksek_acik']}) · kapalı {o['kapali']}" + (f" · kapanma %{o['kapanma']}" if o["kapanma"] is not None else "")
          + (f" · zamanında kapanma %{o['zamaninda_oran']}" if o["zamaninda_oran"] is not None else ""))
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
