import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
GUN = date(2026, 9, 30)


def s(kod, lot, skt, miktar, gunluk=None, fiyat=None):
    return {"kod": kod, "ad": "", "lot": lot, "skt": date(2026, *skt), "miktar": D(miktar), "fiyat": None if fiyat is None else D(fiyat), "depo": "", "gunluk": None if gunluk is None else D(gunluk)}


def veri():
    return [s("M1", "L2", (12, 31), 100, 2, 5), s("M1", "L1", (10, 31), 50, 2, 5), s("M1", "L3", (8, 31), 10, 2, 5), s("M2", "L4", (11, 15), 60, 1, 10), s("m2", "L5", (12, 15), 20, None, 10),
            s("M3", "L6", (11, 1), 5)]


class TakipTesti(unittest.TestCase):
    def test_fefo(self):
        r = main.takip(veri(), GUN)
        self.assertEqual([(x["lot"], x["kalan"], x["durum"], x["kullan"], x["risk"], x["tutar"]) for x in r["lotlar"]],
                         [("L3", -30, "Süresi geçmiş", D(0), D(10), D("50.00")), ("L1", 31, "Yaklaşıyor", D(50), D(0), D("0.00")), ("L6", 32, "Yaklaşıyor", None, None, None),
                          ("L4", 46, "Tükenmeyecek", D(46), D(14), D("140.00")), ("L5", 76, "Tükenmeyecek", D(16), D(4), D("40.00")), ("L2", 92, "Normal", D(100), D(0), D("0.00"))])
        self.assertEqual([(m["kod"], m["lot"], m["miktar"], m["hiz"], m["stok_gun"], m["en_yakin"], m["risk"], m["risk_tutar"]) for m in r["malzemeler"]],
                         [("M2", 2, D(80), D("1.00"), D("80.00"), date(2026, 11, 15), D(18), D("180.00")), ("M1", 3, D(160), D("2.00"), D("75.00"), date(2026, 10, 31), D(10), D("50.00")),
                          ("M3", 1, D(5), None, None, date(2026, 11, 1), D(0), D(0))])
        self.assertEqual((r["ozet"], r["risk_tutar"]), ({"Süresi geçmiş": 1, "Tükenmeyecek": 2, "Yaklaşıyor": 2, "Normal": 1}, D("230.00")))
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Bilgi", "Tüketim bilgisi yok", "M3"), ("Orta", "Süresinde tükenmeyecek", "M2 · L4"), ("Orta", "Süresinde tükenmeyecek", "m2 · L5"), ("Yüksek", "Süresi geçmiş stok", "M1 · L3")])

    def test_secenekler(self):
        r = main.takip(veri(), GUN, yakin=30, tuketim={"m2": D(3), "m3": D("0.1")})
        self.assertEqual([(x["lot"], x["durum"], x["risk"]) for x in r["lotlar"]],
                         [("L3", "Süresi geçmiş", D(10)), ("L1", "Normal", D(0)), ("L6", "Tükenmeyecek", D(2)), ("L4", "Normal", D(0)), ("L5", "Normal", D(0)), ("L2", "Normal", D(0))])
        self.assertEqual([main.skt_oku(x) for x in ("15.11.2026", "10/2026", "2026-12", "2.2027", "13/2026", "x")],
                         [date(2026, 11, 15), date(2026, 10, 31), date(2026, 12, 31), date(2027, 2, 28), None, None])
        with self.assertRaises(ValueError):
            main.takip([], GUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "ilac_stok_lotlari.csv", cikti, GUN)
            self.assertEqual((r["ozet"], r["risk_tutar"]), ({"Süresi geçmiş": 1, "Tükenmeyecek": 4, "Yaklaşıyor": 2, "Normal": 6}, D("15389.50")))
            self.assertEqual([(x["lot"], x["risk"]) for x in r["lotlar"] if x["risk"]], [("OMP1120", D(25)), ("SFT5512", D(25)), ("OND3305", D(44)), ("ENX7781", D(82)), ("AMK9012", D(74))])
            self.assertEqual([(u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Süresi geçmiş stok", "IL-0412 · OMP1120"), ("Süresinde tükenmeyecek", "IL-0215 · SFT5512"), ("Süresinde tükenmeyecek", "IL-0330 · ENX7781"),
                              ("Süresinde tükenmeyecek", "IL-0520 · OND3305"), ("Süresinde tükenmeyecek", "IL-0611 · AMK9012"), ("Tüketim bilgisi yok", "TM-2010")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Lotlar", "Malzemeler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, z = Path(tmp) / "s.csv", Path(tmp) / "t.csv"
            y.write_text("İlaç Kodu;Lot;Miad;Stok Miktarı;Aylık Tüketim\nA;1;12/2026;1.200;300\nA;2;x;5;\nB;3;01.01.2027;0;\nC;4;01.01.2027;-1;\nD;5;01.01.2027;5;y\n;6;01.01.2027;5;\n", encoding="utf-8")
            z.write_text("İlaç Kodu;Çıkış Miktarı\nA;900\na;90\nB;x\n", encoding="utf-8")
            l, uy = main.oku_stok(y)
            self.assertEqual(([(x["kod"], x["skt"], x["miktar"], x["gunluk"]) for x in l], len(uy)), ([("A", date(2026, 12, 31), D(1200), D(10))], 3))
            t, uy = main.oku_tuketim(z, 90)
            self.assertEqual((t, len(uy)), ({"a": D(11)}, 1))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--stok", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tuketim", str(z), "--gun", "90", "--tarih", "30.09.2026", "--yakin", "60"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--gun", "0"], ["--yakin", "-1"])], [2] * 3)
            self.assertEqual(main.main(arg + ["--tuketim", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("İlaç Kodu;Lot\nA;1\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
