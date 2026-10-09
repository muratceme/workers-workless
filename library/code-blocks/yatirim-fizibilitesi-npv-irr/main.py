"""
Yatırım Fizibilitesi (NPV/IRR) — Workers / Workless kod bloğu
Finans › Finansal Analist

Yatırım nakit akışlarından net bugünkü değer, iç verim oranı ve geri dönüş süresini hesaplar:
  - Serbest nakit akışı = gelir − nakit gider − vergi − yatırım harcaması − işletme sermayesi artışı + hurda değeri.
    Vergi = max(0, gelir − nakit gider − amortisman − devreden zarar) × vergi oranı; zarar en fazla 5 yıl taşınır
    (KVK md. 9). Vergi oranı --vergi ile verilir.
  - NBD (NPV) = Σ NA_t ÷ (1 + r)^t. Akışlar sabit fiyatlarla (reel) verildiyse --akis reel ve --enflasyon ile nominal
    iskonto oranı reel orana çevrilir: (1 + nominal) ÷ (1 + enflasyon) − 1.
  - İVO (IRR): NBD = 0 yapan oran (ikiye bölme); işaret birden çok kez değişiyorsa uyarı (birden çok İVO olabilir).
  - Değiştirilmiş İVO (MIRR): negatif akışlar finansman oranıyla bugüne, pozitif akışlar yeniden yatırım oranıyla
    son yıla taşınır.
  - Geri dönüş süresi (basit ve iskontolu, yıl içi doğrusal), kârlılık endeksi = (NBD + BD yatırım) ÷ BD yatırım.
  - Duyarlılık: iskonto oranına göre NBD tablosu ve grafiği.
Rapor: özet, yıllık nakit akışı, duyarlılık, varsayımlar, uyarılar. İnternete bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: 5 yıllık proje, %30 iskonto, %25 vergi
    python main.py --akislar akislar.xlsx --iskonto 35 --vergi 25
    python main.py --akislar akislar.xlsx --iskonto 35 --akis reel --enflasyon 25 --finansman 40 --yeniden-yatirim 30
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
ZARAR_TASIMA_YIL = 5

SUTUNLAR = {"yil": ("yil", "donem", "periyot"), "yatirim": ("yatirim harcamasi", "yatirim", "capex"), "gelir": ("gelir", "satis geliri", "hasilat"),
            "gider": ("nakit gider", "isletme gideri", "gider"), "amortisman": ("amortisman",),
            "is": ("isletme sermayesi degisimi", "isletme sermayesi artisi", "net isletme sermayesi degisimi"), "hurda": ("hurda degeri", "kalinti deger", "hurda"),
            "na": ("nakit akisi", "serbest nakit akisi", "net nakit akisi")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def sayi(x) -> float | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip().replace("TL", "").replace("%", "").replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return float(Decimal(s))
    except InvalidOperation:
        return None


def tl(x: float | None) -> str:
    return "—" if x is None else f"{x:,.0f}".replace(",", ".")


def yuzde(x: float | None) -> str:
    return "—" if x is None else f"%{x * 100:.2f}".replace(".", ",")


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


# ----------------------------------------------------------------------------
# Finansal fonksiyonlar
# ----------------------------------------------------------------------------

def npv(oran: float, akislar: list[float]) -> float:
    return sum(a / (1 + oran) ** t for t, a in enumerate(akislar))


def isaret_degisimi(akislar: list[float]) -> int:
    isaretler = [a > 0 for a in akislar if a != 0]
    return sum(1 for x, y in zip(isaretler, isaretler[1:]) if x != y)


def irr(akislar: list[float], alt: float = -0.99, ust: float = 10.0) -> float | None:
    f_alt, f_ust = npv(alt, akislar), npv(ust, akislar)
    if f_alt * f_ust > 0:
        return None
    for _ in range(200):
        orta = (alt + ust) / 2
        f = npv(orta, akislar)
        if abs(f) < 1e-9 or ust - alt < 1e-12:
            return orta
        if f * f_alt < 0:
            ust = orta
        else:
            alt, f_alt = orta, f
    return (alt + ust) / 2


def mirr(akislar: list[float], finansman: float, yeniden: float) -> float | None:
    n = len(akislar) - 1
    neg = sum(a / (1 + finansman) ** t for t, a in enumerate(akislar) if a < 0)
    poz = sum(a * (1 + yeniden) ** (n - t) for t, a in enumerate(akislar) if a > 0)
    if neg >= 0 or poz <= 0 or n <= 0:
        return None
    return (poz / -neg) ** (1 / n) - 1


def geri_donus(akislar: list[float], oran: float | None = None) -> float | None:
    kum = 0.0
    for t, a in enumerate(akislar):
        d = a / (1 + oran) ** t if oran is not None else a
        once = kum
        kum += d
        if t > 0 and once < 0 <= kum:
            return t - 1 + (-once / d if d else 0)
    return None


# ----------------------------------------------------------------------------
# Nakit akışı
# ----------------------------------------------------------------------------

@dataclass
class Yil:
    yil: int
    yatirim: float
    gelir: float
    gider: float
    amortisman: float
    isd: float
    hurda: float
    na_verilen: float | None
    matrah: float = 0.0
    mahsup: float = 0.0
    vergi: float = 0.0

    @property
    def faaliyet(self) -> float:
        return self.gelir - self.gider

    @property
    def na(self) -> float:
        if self.na_verilen is not None:
            return self.na_verilen
        return self.faaliyet - self.vergi - self.yatirim - self.isd + self.hurda


def vergi_hesapla(yillar: list[Yil], oran: float) -> None:
    zararlar: list[list] = []        # [yıl, kalan zarar]
    for y in yillar:
        m = y.faaliyet - y.amortisman
        zararlar = [z for z in zararlar if y.yil - z[0] <= ZARAR_TASIMA_YIL and z[1] > 0]
        if m < 0:
            zararlar.append([y.yil, -m])
            y.matrah, y.mahsup, y.vergi = m, 0.0, 0.0
            continue
        mahsup = 0.0
        for z in zararlar:
            al = min(z[1], m - mahsup)
            z[1] -= al
            mahsup += al
        y.matrah, y.mahsup = m, mahsup
        y.vergi = (m - mahsup) * oran


def analiz_et(yillar: list[Yil], iskonto: float, vergi: float, akis: str, enflasyon: float, finansman: float | None, yeniden: float | None) -> dict:
    uy = []
    if not any(y.na_verilen is not None for y in yillar):
        vergi_hesapla(yillar, vergi)
        if vergi and not any(y.amortisman for y in yillar) and any(y.yatirim for y in yillar):
            uy.append({"onem": "Orta", "tur": "Amortisman yok", "aciklama": "Vergi hesaplanıyor ama amortisman verilmedi; vergi fazla, NBD düşük çıkar"})
    oran = (1 + iskonto) / (1 + enflasyon) - 1 if akis == "reel" else iskonto
    na = [y.na for y in yillar]
    s = {"yillar": yillar, "na": na, "oran": oran, "iskonto": iskonto, "akis": akis, "enflasyon": enflasyon, "vergi": vergi,
         "npv": npv(oran, na), "irr": irr(na), "mirr": None, "gd": geri_donus(na), "igd": geri_donus(na, oran), "uyarilar": uy}
    s["mirr"] = mirr(na, finansman if finansman is not None else oran, yeniden if yeniden is not None else oran)
    bd_yatirim = sum(-a / (1 + oran) ** t for t, a in enumerate(na) if a < 0)
    s["bd_yatirim"] = bd_yatirim
    s["pi"] = (s["npv"] + bd_yatirim) / bd_yatirim if bd_yatirim else None
    if isaret_degisimi(na) > 1:
        uy.append({"onem": "Orta", "tur": "Birden çok işaret değişimi", "aciklama": "Nakit akışının işareti birden çok kez değişiyor; birden çok İVO olabilir, "
                   "kararı NBD ve Değiştirilmiş İVO ile verin"})
    if s["irr"] is None:
        uy.append({"onem": "Orta", "tur": "İVO bulunamadı", "aciklama": "−%99 ile %1000 arasında NBD'yi sıfırlayan oran yok"})
    if s["gd"] is None:
        uy.append({"onem": "Yüksek", "tur": "Geri dönmüyor", "aciklama": "Proje süresi içinde kümülatif nakit akışı pozitife dönmüyor"})
    elif s["igd"] is None:
        uy.append({"onem": "Orta", "tur": "İskontolu geri dönüş yok", "aciklama": "İskonto oranıyla proje süresinde geri dönmüyor"})
    if s["npv"] < 0:
        uy.append({"onem": "Yüksek", "tur": "Negatif NBD", "aciklama": f"NBD {tl(s['npv'])} TL; iskonto oranı ({yuzde(oran)}) ile proje değer yaratmıyor"})
    if akis == "reel" and not enflasyon:
        uy.append({"onem": "Orta", "tur": "Enflasyon yok", "aciklama": "Akışlar reel işaretli ama enflasyon verilmedi; nominal oran reel gibi kullanıldı"})
    if akis == "nominal" and iskonto < 0.15:
        uy.append({"onem": "Bilgi", "tur": "İskonto oranı", "aciklama": f"Nominal akış için iskonto oranı {yuzde(iskonto)}; nominal oranın enflasyon "
                   "beklentisini içerdiğinden emin olun (reel akış ise --akis reel)"})
    son = yillar[-1]
    if son.hurda == 0 and not any(y.isd < 0 for y in yillar) and any(y.isd > 0 for y in yillar):
        uy.append({"onem": "Bilgi", "tur": "İşletme sermayesi", "aciklama": "İşletme sermayesi artışları proje sonunda geri alınmamış görünüyor"})
    # Duyarlılık
    adimlar = sorted({round(oran + d, 4) for d in (-0.15, -0.10, -0.05, -0.025, 0, 0.025, 0.05, 0.10, 0.15) if oran + d > -0.99})
    s["duyarlilik"] = [(r, npv(r, na)) for r in adimlar]
    return s


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
UST = Alignment(vertical="top", wrap_text=True)
PF = "#,##0;[Red]-#,##0"


def _baslik(ws, basliklar, genislik=()):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def rapor_yaz(cikti: Path, s: dict) -> None:
    wb = Workbook()
    oz = wb.active
    oz.title = "Özet"
    _baslik(oz, ["Gösterge", "Değer", "Açıklama"], (34, 18, 70))
    karar = "Kabul edilebilir (NBD > 0)" if s["npv"] > 0 else "Reddedilmeli (NBD < 0)" if s["npv"] < 0 else "Başa baş"
    for ad, v, fmt, a in [("Net bugünkü değer (NBD)", s["npv"], PF, f"İskonto oranı {yuzde(s['oran'])}" + (" (reel)" if s["akis"] == "reel" else "")),
                          ("İç verim oranı (İVO)", s["irr"], "0.00%", "NBD'yi sıfırlayan oran; iskonto oranından büyükse proje değer yaratır"),
                          ("Değiştirilmiş İVO (MIRR)", s["mirr"], "0.00%", "Finansman ve yeniden yatırım oranlarıyla"),
                          ("Kârlılık endeksi", s["pi"], "0.00", "(NBD + bugünkü değerle yatırım) ÷ bugünkü değerle yatırım; > 1 olumlu"),
                          ("Basit geri dönüş süresi (yıl)", s["gd"], "0.00", "Kümülatif nakit akışının pozitife döndüğü an (yıl içi doğrusal)"),
                          ("İskontolu geri dönüş süresi (yıl)", s["igd"], "0.00", "İskontolu kümülatif akışla"),
                          ("Bugünkü değerle yatırım", s["bd_yatirim"], PF, "Negatif akışların bugünkü değeri"),
                          ("Değerlendirme", karar, None, "Tek başına karar ölçütü değildir; riskler ve senaryolar birlikte değerlendirilmelidir")]:
        oz.append([ad, v if v is not None else "—", a])
        if fmt and v is not None:
            oz.cell(oz.max_row, 2).number_format = fmt
    oz.cell(2, 2).fill = PatternFill("solid", fgColor="E3F4E1" if s["npv"] > 0 else "FDE2E1")

    na = wb.create_sheet("Nakit Akışı")
    _baslik(na, ["Yıl", "Gelir", "Nakit Gider", "Faaliyet Nakdi", "Amortisman", "Vergi Matrahı", "Zarar Mahsubu", "Vergi", "Yatırım", "İşletme Sermayesi Artışı",
                 "Hurda Değeri", "Serbest Nakit Akışı", "İskonto Çarpanı", "Bugünkü Değer", "Kümülatif", "İskontolu Kümülatif"],
            (6, 13, 13, 13, 12, 13, 12, 12, 13, 13, 12, 15, 10, 14, 14, 14))
    kum = ikum = 0.0
    for t, y in enumerate(s["yillar"]):
        c = 1 / (1 + s["oran"]) ** t
        kum += y.na
        ikum += y.na * c
        na.append([y.yil, y.gelir, y.gider, y.faaliyet, y.amortisman, y.matrah, y.mahsup, y.vergi, y.yatirim, y.isd, y.hurda, y.na, c, y.na * c, kum, ikum])
        for j in range(2, 17):
            na.cell(na.max_row, j).number_format = "0.0000" if j == 13 else PF
    na.append(["Toplam", *[None] * 10, sum(s["na"]), None, s["npv"]])
    for j in (12, 14):
        na.cell(na.max_row, j).number_format = PF
        na.cell(na.max_row, j).font = Font(bold=True)
    na.freeze_panes = "B2"

    dy = wb.create_sheet("Duyarlılık")
    _baslik(dy, ["İskonto Oranı", "NBD"], (14, 18))
    for r, v in s["duyarlilik"]:
        dy.append([r, v])
        dy.cell(dy.max_row, 1).number_format = "0.0%"
        dy.cell(dy.max_row, 2).number_format = PF
        if abs(r - s["oran"]) < 1e-9:
            dy.cell(dy.max_row, 1).font = dy.cell(dy.max_row, 2).font = Font(bold=True)
    g = LineChart()
    g.title, g.height, g.width = "İskonto oranına göre NBD", 8, 16
    g.add_data(Reference(dy, min_col=2, min_row=1, max_row=dy.max_row), titles_from_data=True)
    g.set_categories(Reference(dy, min_col=1, min_row=2, max_row=dy.max_row))
    g.legend = None
    dy.add_chart(g, "D2")

    vs = wb.create_sheet("Varsayımlar")
    _baslik(vs, ["Varsayım", "Değer"], (36, 70))
    for a, v in [("İskonto oranı (girilen)", yuzde(s["iskonto"])), ("Akış türü", "Reel (sabit fiyatlarla)" if s["akis"] == "reel" else "Nominal (cari fiyatlarla)"),
                 ("Enflasyon (reel akışta)", yuzde(s["enflasyon"]) if s["akis"] == "reel" else "—"), ("Kullanılan iskonto oranı", yuzde(s["oran"])),
                 ("Vergi oranı", yuzde(s["vergi"])), ("Zarar mahsubu", f"En fazla {ZARAR_TASIMA_YIL} yıl (KVK md. 9)"),
                 ("Zamanlama", "Akışlar yıl sonunda gerçekleşir; 0. yıl bugündür"),
                 ("Kapsam dışı", "Finansman giderleri (iskonto oranında yansır), yatırım teşvikleri ve vergi istisnaları, hurda satış kazancının vergisi")]:
        vs.append([a, v])
        vs.cell(vs.max_row, 2).alignment = UST

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Açıklama"], (9, 26, 100))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["aciklama"]])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 3).alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(yol: Path, cikti: Path, iskonto: float, vergi: float = 0.25, akis: str = "nominal", enflasyon: float = 0.0,
             finansman: float | None = None, yeniden: float | None = None) -> dict:
    yillar = []
    for r in kayitlar(yol, SUTUNLAR, ("yil",)):
        y = sayi(r.get("yil"))
        if y is None:
            continue
        al = lambda k: sayi(r.get(k)) or 0.0  # noqa: E731
        yillar.append(Yil(int(y), abs(al("yatirim")), al("gelir"), abs(al("gider")), abs(al("amortisman")), al("is"), al("hurda"), sayi(r.get("na"))))
    if len(yillar) < 2:
        raise ValueError(f"{yol.name}: en az iki yıllık nakit akışı gerekli")
    yillar.sort(key=lambda y: y.yil)
    if [y.yil for y in yillar] != list(range(yillar[0].yil, yillar[0].yil + len(yillar))):
        raise ValueError("Yıllar ardışık olmalı (0, 1, 2, … veya 2027, 2028, …)")
    s = analiz_et(yillar, iskonto, vergi, akis, enflasyon, finansman, yeniden)
    rapor_yaz(cikti, s)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Yatırım nakit akışlarından NBD, İVO, MIRR, geri dönüş süresi ve duyarlılık tablosu hesaplar.")
    p.add_argument("--akislar", type=Path, default=ORNEK / "nakit_akislari.csv",
                   help="Yıl, Yatırım Harcaması, Gelir, Nakit Gider, Amortisman, İşletme Sermayesi Değişimi, Hurda Değeri (veya doğrudan Nakit Akışı)")
    p.add_argument("--iskonto", type=float, default=30, help="İskonto oranı / sermaye maliyeti %% (varsayılan 30)")
    p.add_argument("--vergi", type=float, default=25, help="Kurumlar vergisi oranı %% (varsayılan 25; güncel oranı kontrol edin)")
    p.add_argument("--akis", choices=("nominal", "reel"), default="nominal", help="Akışlar cari fiyatlarla mı (nominal) sabit fiyatlarla mı (reel)")
    p.add_argument("--enflasyon", type=float, default=0, help="Reel akışta beklenen yıllık enflasyon %%")
    p.add_argument("--finansman", type=float, help="MIRR finansman oranı %% (varsayılan iskonto oranı)")
    p.add_argument("--yeniden-yatirim", type=float, help="MIRR yeniden yatırım oranı %% (varsayılan iskonto oranı)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "yatirim_fizibilitesi.xlsx")
    a = p.parse_args(argv)
    if not a.akislar.exists():
        print(f"[X] Dosya bulunamadı: {a.akislar}")
        return 1
    if a.iskonto <= -99:
        print("[X] --iskonto −99'dan büyük olmalı")
        return 2
    try:
        s = calistir(a.akislar, a.cikti, a.iskonto / 100, a.vergi / 100, a.akis, a.enflasyon / 100,
                     None if a.finansman is None else a.finansman / 100, None if a.yeniden_yatirim is None else a.yeniden_yatirim / 100)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    yil = lambda x: "—" if x is None else f"{x:.2f}".replace(".", ",")  # noqa: E731
    print(f"[OK] NBD {tl(s['npv'])} TL (iskonto {yuzde(s['oran'])}) · İVO {yuzde(s['irr'])} · MIRR {yuzde(s['mirr'])} · "
          f"geri dönüş {yil(s['gd'])} yıl · iskontolu {yil(s['igd'])} yıl")
    for u in s["uyarilar"]:
        if u["onem"] != "Bilgi":
            print(f"[!] {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
