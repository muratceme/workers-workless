import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


class FonksiyonTesti(unittest.TestCase):
    def test_bilinen_ornekler(self):
        self.assertAlmostEqual(main.irr([-123400, 36200, 54800, 48100]), 0.0596, places=4)      # Wikipedia "Internal rate of return" örneği
        self.assertAlmostEqual(main.npv(0.10, [-1000, 1100]), 0.0, places=9)
        self.assertAlmostEqual(main.npv(0.10, [-100, 60, 60]), -100 + 60 / 1.1 + 60 / 1.21, places=9)
        self.assertAlmostEqual(main.mirr([-1000, 1210], 0.10, 0.10), 0.21, places=9)
        self.assertAlmostEqual(main.geri_donus([-100, 40, 40, 40]), 2.5)
        self.assertIsNone(main.irr([100, 100]))
        self.assertEqual(main.isaret_degisimi([-100, 230, -132]), 2)

    def test_vergi_ve_zarar_mahsubu(self):
        Y = main.Yil
        yillar = [Y(0, 1000, 0, 0, 0, 0, 0, None), Y(1, 0, 100, 150, 50, 0, 0, None), Y(2, 0, 400, 100, 50, 0, 0, None)]
        main.vergi_hesapla(yillar, 0.25)
        self.assertEqual((yillar[1].matrah, yillar[1].vergi), (-100, 0))
        self.assertEqual((yillar[2].mahsup, yillar[2].vergi), (100, 37.5))         # (250 − 100) × %25


class OrnekVeriTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "y.xlsx"
            s = main.calistir(ORNEK / "nakit_akislari.csv", cikti, 0.30)
            self.assertEqual(s["na"], [-13500000, 2775000, 3725000, 4275000, 4650000, 8650000])
            self.assertAlmostEqual(main.npv(s["irr"], s["na"]), 0, delta=1e-3)
            self.assertAlmostEqual(s["gd"], 3 + 2725000 / 4650000)
            self.assertLess(s["npv"], 0)
            self.assertIn("Negatif NBD", {u["tur"] for u in s["uyarilar"]})
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Özet", "Nakit Akışı", "Duyarlılık", "Varsayımlar", "Uyarılar"])
            s2 = main.calistir(ORNEK / "nakit_akislari.csv", cikti, 0.30, akis="reel", enflasyon=0.20)
            self.assertAlmostEqual(s2["oran"], 1.30 / 1.20 - 1)
            self.assertGreater(s2["npv"], 0)

    def test_dogrudan_nakit_akisi_ve_hata(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t / "a.csv").write_text("Yıl;Nakit Akışı\n0;-100\n1;230\n2;-132\n", encoding="utf-8")
            s = main.calistir(t / "a.csv", t / "o.xlsx", 0.10, vergi=0)
            self.assertIn("Birden çok işaret değişimi", {u["tur"] for u in s["uyarilar"]})
            (t / "b.csv").write_text("Yıl;Nakit Akışı\n0;-100\n2;50\n", encoding="utf-8")
            with self.assertRaises(ValueError):
                main.calistir(t / "b.csv", t / "o.xlsx", 0.10)

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx"), "--iskonto", "15"]), 0)
            self.assertEqual(main.main(["--iskonto", "-100"]), 2)
            self.assertEqual(main.main(["--akislar", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
