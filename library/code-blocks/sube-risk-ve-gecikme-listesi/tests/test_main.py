import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
K, Y, O, E = "1 – Kritik", "2 – Yüksek", "3 – Orta", "4 – Erken uyarı"


def k(no, gun=0, geciken=None, risk=1000, uyari="", py="P", kredi="K"):
    return {"no": no, "ad": "Ad", "kredi": kredi, "urun": "Ü", "risk": D(str(risk)), "geciken": D(str(geciken if geciken is not None else (10 if gun else 0))), "gun": gun, "uyari": uyari, "py": py}


class DegerlendirTesti(unittest.TestCase):
    def test_dilim_ve_oncelik(self):
        gunler = (0, 1, 30, 31, 60, 61, 90, 91)
        r = main.degerlendir([k(str(g), g) for g in gunler])
        self.assertEqual({x["no"]: x["dilim"] for x in r["krediler"]}, {"0": "Gecikmesiz", "1": "1–30", "30": "1–30", "31": "31–60", "60": "31–60", "61": "61–90", "90": "61–90", "91": "91+"})
        self.assertEqual([(x["no"], x["oncelik"], x["takibe_kalan"]) for x in r["musteriler"]], [("91", K, -1), ("90", K, 0), ("61", K, 29), ("60", Y, 30), ("31", Y, 59), ("30", O, 60), ("1", O, 89)])
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]), [("Takip eşiği aşıldı", "91 · Ad"), ("Takip eşiğine yaklaşıyor", "61 · Ad"), ("Takip eşiğine yaklaşıyor", "90 · Ad")])
        self.assertEqual([(d["dilim"], d["kredi"], d["risk"], d["pay"], d["geciken"]) for d in r["dilimler"]],
                         [("Gecikmesiz", 1, D(1000), D("12.50"), D(0)), ("1–30", 2, D(2000), D("25.00"), D(20)), ("31–60", 2, D(2000), D("25.00"), D(20)), ("61–90", 2, D(2000), D("25.00"), D(20)), ("91+", 1, D(1000), D("12.50"), D(10))])
        self.assertEqual((r["toplam"], r["say"]), ({"risk": D(8000), "gecikmeli_risk": D(7000), "geciken": D(70), "oran": D("87.50")}, {K: 3, Y: 2, O: 2, E: 0}))
        r = main.degerlendir([k("a", 16), k("b", 31), k("c", 45)], dilimler=(15, 30, 60), takip=45)
        self.assertEqual(([(x["no"], x["oncelik"], x["takibe_kalan"]) for x in r["musteriler"]], [d["dilim"] for d in r["dilimler"]], sorted(u["tur"] for u in r["uyarilar"])),
                         ([("c", K, 0), ("b", K, 14), ("a", Y, 29)], ["Gecikmesiz", "1–15", "16–30", "31–60", "61+"], ["Takip eşiğine yaklaşıyor"] * 2))
        self.assertEqual([x["oncelik"] for x in main.degerlendir([k("a", 31), k("b", 5)], dilimler=(30,))["musteriler"]], [K, O])
        with self.assertRaises(ValueError):
            main.degerlendir([])

    def test_musteri_ve_ozet(self):
        l = [k("M1", 72, 85, 1200, py="Z", kredi="K1"), k("m1", 0, 0, 400, "limit dolu", py="Z", kredi="K2"), k("M2", 20, 250000, 900000, py="O"), k("M3", 20, "249999.99", 900000, py="O"), k("M4", 0, 0, 500, "Çek iadesi", py="o"),
             k("M4", 0, 0, 100, "Çek iadesi", py="O"), k("M5", py="O"), k("M6", 5, 0, 100, py="Z"), k("M7", 0, 50, 10, py="Z")]
        r = main.degerlendir(l)
        self.assertEqual([(x["no"], x["oncelik"], x["kredi"], x["geciken_kredi"], x["risk"], x["geciken"], x["gun"], x["uyari"], x["takibe_kalan"]) for x in r["musteriler"]],
                         [("M1", K, 2, 1, D(1600), D(85), 72, "limit dolu", 18), ("M2", Y, 1, 1, D(900000), D(250000), 20, "", 70), ("M3", O, 1, 1, D(900000), D("249999.99"), 20, "", 70), ("M6", O, 1, 1, D(100), D(0), 5, "", 85),
                          ("M4", E, 2, 0, D(600), D(0), 0, "Çek iadesi", None)])
        self.assertEqual(r["yoneticiler"], [{"py": "O", "risk": D(1801600), "geciken": D("499999.99"), "oran": D("99.91"), "musteri": 3, "kritik": 0}, {"py": "Z", "risk": D(1710), "geciken": D(135), "oran": D("76.02"), "musteri": 2, "kritik": 1}])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         sorted([("Yüksek", "Takip eşiğine yaklaşıyor", "M1 · Ad"), ("Bilgi", "Erken uyarı", "M4 · Ad"), ("Orta", "Gecikme bilgisi tutarsız", "M6 · Ad · K"), ("Orta", "Gecikme bilgisi tutarsız", "M7 · Ad · K"),
                                 ("Orta", "Geciken tutar riskten büyük", "M7 · Ad · K")]))
        self.assertEqual(main.degerlendir(l, buyuk=D(100000))["musteriler"][2]["oncelik"], Y)
        self.assertEqual(main.degerlendir([k("a", risk=0)])["toplam"]["oran"], None)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "krediler.csv", cikti)
            self.assertEqual([(x["no"], x["ad"], x["oncelik"], x["risk"], x["geciken"], x["gun"], x["takibe_kalan"]) for x in r["musteriler"]],
                             [("3006", "Zeta Mobilya", K, D(350000), D(350000), 104, -14), ("3001", "Alfa Market", K, D(1600000), D(85000), 72, 18), ("3002", "Beta Yapı", Y, D(2500000), D(310000), 45, 45),
                              ("3004", "Delta Tekstil", Y, D(900000), D(300000), 20, 70), ("3003", "Gama Gıda", O, D(800000), D(57000), 30, 60), ("3005", "Epsilon Otomotiv", E, D(1500000), D(0), 0, None)])
            self.assertEqual([(d["dilim"], d["kredi"], d["risk"], d["geciken"]) for d in r["dilimler"]],
                             [("Gecikmesiz", 3, D(2700000), D(0)), ("1–30", 3, D(1700000), D(357000)), ("31–60", 1, D(2500000), D(310000)), ("61–90", 1, D(1200000), D(85000)), ("91+", 1, D(350000), D(350000))])
            self.assertEqual((r["toplam"], r["say"]), ({"risk": D(8450000), "gecikmeli_risk": D(5750000), "geciken": D(1102000), "oran": D("68.05")}, {K: 2, Y: 2, O: 1, E: 1}))
            self.assertEqual(r["yoneticiler"], [{"py": "Okan Er", "risk": D(3550000), "geciken": D(707000), "oran": D("57.75"), "musteri": 4, "kritik": 1},
                                                {"py": "Zeynep Ak", "risk": D(4900000), "geciken": D(395000), "oran": D("75.51"), "musteri": 2, "kritik": 1}])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Yüksek", "Takip eşiği aşıldı", "3006 · Zeta Mobilya"), ("Yüksek", "Takip eşiğine yaklaşıyor", "3001 · Alfa Market"), ("Bilgi", "Erken uyarı", "3005 · Epsilon Otomotiv")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Öncelik Listesi", "Özet", "Krediler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "k.csv"
            y.write_text("Müşteri Numarası;Unvan;Hesap No;Kredi Türü;Bakiye;Gecikmiş Tutar;Gecikme;Sinyal;Portföy\n7;A;H1;Rotatif;1.250.000,50;12.500;45;;\n8;B;;;5;;;x;Can\n9;C;;;x;;;;\n10;C;;;5;-1;;;\n11;C;;;5;;1,5;;\n;C;;;5;;;;\n",
                         encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["kredi"], x["urun"], x["risk"], x["geciken"], x["gun"], x["uyari"], x["py"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("7", "H1", "Rotatif", D("1250000.50"), D(12500), 45, "", "(atanmamış)"), ("8", "", "", D(5), D(0), 0, "x", "Can")], [("Kredi satırı okunamadı", f"{n} · C") for n in (9, 10, 11)]))
            self.assertEqual([main.dilim_coz(x) for x in ("30;60;90", "15", "60;30", "0;30", "x")], [(30, 60, 90), (15,), None, None, None])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--krediler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--dilimler", "15;30;60;90", "--takip", "90", "--buyuk", "500.000"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--dilimler", "60;30"], ["--takip", "0"], ["--buyuk", "0"], ["--buyuk", "x"])], [2] * 4)
            self.assertEqual(main.main(["--krediler", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Müşteri No;Ürün\nA;B\n", "Müşteri No;Risk\n;1\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
