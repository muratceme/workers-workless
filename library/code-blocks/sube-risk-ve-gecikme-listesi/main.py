"""
Şube Risk ve Gecikme Listesi — Gerisi Bende kod bloğu
Şube Bankacılığı › Şube Müdürü

Şube portföyündeki gecikmeye düşen ve erken uyarı veren kredileri tutar ve gün bazında önceliklendirir.
  Gecikme dilimi       = Gecikmesiz · 1–30 · 31–60 · 61–90 · 91+ gün   (--dilimler ile değişir)
  Müşteri özeti        = müşterinin kredilerinin risk ve geciken tutar toplamı; en uzun gecikme günü
  Öncelik              = 1 – Kritik      : en uzun gecikme ikinci dilim sınırını (60 gün) aşmış
                         2 – Yüksek      : en uzun gecikme ilk dilim sınırını (30 gün) aşmış ya da geciken tutar --buyuk ve üstünde
                         3 – Orta        : gecikme var
                         4 – Erken uyarı : gecikme yok, erken uyarı notu var
  Sıra                 = öncelik → gecikme günü (uzundan kısaya) → geciken tutar (büyükten küçüğe)
  Takip eşiğine kalan  = --takip (90) − en uzun gecikme günü
  Gecikme oranı %      = gecikmedeki kredilerin risk bakiyesi ÷ toplam risk bakiyesi × 100
Öncelik kuralı bir çalışma varsayımıdır. Kredi sınıflandırması ve karşılık hesabı bu paketin konusu değildir.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
SUTUN = {"no": ("musteri no", "musteri numarasi"), "ad": ("musteri", "musteri adi", "unvan"), "kredi": ("kredi no", "hesap no", "kredi numarasi"), "urun": ("urun", "kredi turu", "urun turu"),
         "risk": ("risk bakiyesi", "risk", "bakiye", "kalan anapara"), "geciken": ("geciken tutar", "gecikmis tutar", "geciken"), "gun": ("gecikme gun", "gecikme gunu", "gecikme gun sayisi", "gecikme"),
         "uyari": ("erken uyari", "erken uyari sinyali", "sinyal"), "py": ("portfoy yoneticisi", "musteri temsilcisi", "portfoy")}
ONCELIK_RENK = {"1 – Kritik": KIRMIZI, "2 – Yüksek": SARI, "3 – Orta": MAVI, "4 – Erken uyarı": None}
Z, YUZ = Decimal(0), Decimal(100)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("no", "risk")):
        no = metin(r["no"])
        if not no:
            continue
        kim = f"{no} · {metin(r['ad'])}".strip(" ·")
        risk, geciken, gun = para(r["risk"]), para(r["geciken"]) if metin(r["geciken"]) else Z, para(r["gun"]) if metin(r["gun"]) else Z
        if risk is None or risk < 0 or geciken is None or geciken < 0 or gun is None or gun < 0 or gun != gun.to_integral_value():
            uy.append(uyari("Yüksek", "Kredi satırı okunamadı", kim, f"Satır {r['_satir']}: risk '{metin(r['risk'])}', geciken tutar '{metin(r['geciken'])}', gecikme gün '{metin(r['gun'])}'. Satır alınmadı."))
            continue
        l.append({"no": no, "ad": metin(r["ad"]), "kredi": metin(r["kredi"]), "urun": metin(r["urun"]), "risk": risk, "geciken": geciken, "gun": int(gun), "uyari": metin(r["uyari"]), "py": metin(r["py"]) or "(atanmamış)"})
    return l, uy


def degerlendir(krediler: list[dict], dilimler: tuple[int, ...] = (30, 60, 90), takip: int = 90, buyuk: Decimal = Decimal(250000)) -> dict:
    if not krediler:
        raise ValueError("Değerlendirilecek kredi yok.")
    uy = []
    etiketler = ["Gecikmesiz"] + [f"{a}–{b}" for a, b in zip((1,) + tuple(d + 1 for d in dilimler), dilimler)] + [f"{dilimler[-1] + 1}+"]
    kritik_gun = dilimler[1] if len(dilimler) > 1 else dilimler[0]
    musteriler = {}
    for k in krediler:
        k["dilim"] = etiketler[0] if k["gun"] == 0 else etiketler[1 + next((i for i, d in enumerate(dilimler) if k["gun"] <= d), len(dilimler))]
        kim = f"{k['no']} · {k['ad']} · {k['kredi']}".replace(" ·  ·", " ·").strip(" ·")
        if k["geciken"] > k["risk"]:
            uy.append(uyari("Orta", "Geciken tutar riskten büyük", kim, f"Geciken tutar {k['geciken']:,.2f}, risk bakiyesi {k['risk']:,.2f}. Veriyi kontrol edin."))
        if (k["gun"] > 0) != (k["geciken"] > 0):
            uy.append(uyari("Orta", "Gecikme bilgisi tutarsız", kim, f"Gecikme gün {k['gun']}, geciken tutar {k['geciken']:,.2f}; biri sıfır, diğeri değil."))
        m = musteriler.setdefault(katla(k["no"]), {"no": k["no"], "ad": k["ad"], "py": k["py"], "krediler": []})
        m["krediler"].append(k)
    liste = []
    for m in musteriler.values():
        ks = m["krediler"]
        o = {"no": m["no"], "ad": m["ad"], "py": m["py"], "kredi": len(ks), "geciken_kredi": sum(1 for k in ks if k["gun"] > 0), "risk": sum((k["risk"] for k in ks), Z), "geciken": sum((k["geciken"] for k in ks), Z), "gun": max(k["gun"] for k in ks),
             "uyari": "; ".join(dict.fromkeys(k["uyari"] for k in ks if k["uyari"]))}
        kim = f"{o['no']} · {o['ad']}".strip(" ·")
        if o["gun"] > kritik_gun:
            o["oncelik"] = "1 – Kritik"
        elif o["gun"] > dilimler[0] or (o["gun"] > 0 and o["geciken"] >= buyuk):
            o["oncelik"] = "2 – Yüksek"
        elif o["gun"] > 0:
            o["oncelik"] = "3 – Orta"
        elif o["uyari"]:
            o["oncelik"] = "4 – Erken uyarı"
            uy.append(uyari("Bilgi", "Erken uyarı", kim, f"Gecikme yok; risk {o['risk']:,.2f}. Not: {o['uyari']}"))
        else:
            continue
        o["takibe_kalan"] = takip - o["gun"] if o["gun"] > 0 else None
        if o["gun"] > takip:
            uy.append(uyari("Yüksek", "Takip eşiği aşıldı", kim, f"En uzun gecikme {o['gun']} gün (eşik {takip}). Risk {o['risk']:,.2f}, geciken {o['geciken']:,.2f}. Sınıflandırma ve yasal takip sürecini kredi izleme birimiyle değerlendirin."))
        elif o["oncelik"] == "1 – Kritik":
            uy.append(uyari("Yüksek", "Takip eşiğine yaklaşıyor", kim, f"En uzun gecikme {o['gun']} gün; {takip} günlük eşiğe {o['takibe_kalan']} gün kaldı. Risk {o['risk']:,.2f}, geciken {o['geciken']:,.2f}."))
        liste.append(o)
    liste.sort(key=lambda o: (o["oncelik"], -o["gun"], -o["geciken"], katla(o["no"])))
    toplam, gecikmeli = sum((k["risk"] for k in krediler), Z), sum((k["risk"] for k in krediler if k["gun"] > 0), Z)
    dilim_ozet = [{"dilim": e, "kredi": sum(1 for k in krediler if k["dilim"] == e), "risk": sum((k["risk"] for k in krediler if k["dilim"] == e), Z), "geciken": sum((k["geciken"] for k in krediler if k["dilim"] == e), Z)} for e in etiketler]
    for d in dilim_ozet:
        d["pay"] = yuvarla(d["risk"] * YUZ / toplam) if toplam else None
    listede = {katla(o["no"]): o for o in liste}
    py = {}
    for k in krediler:
        g = py.setdefault(katla(k["py"]), {"py": k["py"], "risk": Z, "geciken": Z, "gecikmeli_risk": Z, "musteriler": set()})
        g["risk"] += k["risk"]
        g["geciken"] += k["geciken"]
        g["gecikmeli_risk"] += k["risk"] if k["gun"] > 0 else Z
        if katla(k["no"]) in listede:
            g["musteriler"].add(katla(k["no"]))
    yoneticiler = [{"py": g["py"], "risk": g["risk"], "geciken": g["geciken"], "oran": yuvarla(g["gecikmeli_risk"] * YUZ / g["risk"]) if g["risk"] else None, "musteri": len(g["musteriler"]),
                    "kritik": sum(1 for n in g["musteriler"] if listede[n]["oncelik"] == "1 – Kritik")} for _, g in sorted(py.items())]
    return {"musteriler": liste, "krediler": sorted(krediler, key=lambda k: (-k["gun"], -k["geciken"], katla(k["no"]))), "dilimler": dilim_ozet, "yoneticiler": yoneticiler,
            "toplam": {"risk": toplam, "gecikmeli_risk": gecikmeli, "geciken": sum((k["geciken"] for k in krediler), Z), "oran": yuvarla(gecikmeli * YUZ / toplam) if toplam else None},
            "say": {o: sum(1 for x in liste if x["oncelik"] == o) for o in ONCELIK_RENK}, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, takip: int) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Öncelik Listesi"
    _baslik(ws, ["Öncelik", "Müşteri No", "Müşteri", "Risk Bakiyesi", "Geciken Tutar", "En Uzun Gecikme (gün)", f"{takip} Güne Kalan", "Kredi", "Geciken Kredi", "Erken Uyarı", "Portföy Yöneticisi", "Aksiyon (siz doldurun)"],
            (16, 12, 28, 16, 16, 11, 10, 7, 9, 40, 22, 36))
    for o in r["musteriler"]:
        _satir(ws, [o["oncelik"], o["no"], o["ad"], o["risk"], o["geciken"] or None, o["gun"] or None, o["takibe_kalan"], o["kredi"], o["geciken_kredi"] or None, o["uyari"], o["py"], None], para_s=(4, 5),
               dolgu={1: ONCELIK_RENK[o["oncelik"]], 7: KIRMIZI if o["takibe_kalan"] is not None and o["takibe_kalan"] < 0 else None}, sar=(10,))
    ws = wb.create_sheet("Özet")
    _baslik(ws, ["Gecikme Dilimi (gün)", "Kredi", "Risk Bakiyesi", "Pay %", "Geciken Tutar"], (22, 8, 18, 9, 18))
    for d in r["dilimler"]:
        _satir(ws, [d["dilim"], d["kredi"], d["risk"], d["pay"], d["geciken"] or None], para_s=(3, 4, 5))
    t = r["toplam"]
    n = _satir(ws, ["TOPLAM", len(r["krediler"]), t["risk"], None, t["geciken"]], para_s=(3, 5))
    for h in ws[n]:
        h.font = Font(bold=True)
    _satir(ws, ["Gecikmedeki kredilerin riski", None, t["gecikmeli_risk"], t["oran"], None], para_s=(3, 4))
    ws.append([])
    _baslik(ws, ["Portföy Yöneticisi", "Listedeki Müşteri", "Risk Bakiyesi", "Gecikme Oranı %", "Geciken Tutar", "Kritik Müşteri"], dondur="A2")
    for y in r["yoneticiler"]:
        _satir(ws, [y["py"], y["musteri"], y["risk"], y["oran"], y["geciken"] or None, y["kritik"] or None], para_s=(3, 4, 5))
    ws.column_dimensions["F"].width = 14
    ws = wb.create_sheet("Krediler")
    _baslik(ws, ["Müşteri No", "Müşteri", "Kredi No", "Ürün", "Risk Bakiyesi", "Geciken Tutar", "Gecikme (gün)", "Dilim", "Erken Uyarı", "Portföy Yöneticisi"], (12, 28, 14, 20, 16, 16, 10, 12, 40, 22))
    for k in r["krediler"]:
        _satir(ws, [k["no"], k["ad"], k["kredi"], k["urun"], k["risk"], k["geciken"] or None, k["gun"] or None, k["dilim"], k["uyari"], k["py"]], para_s=(5, 6), sar=(9,))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, dilimler: tuple[int, ...] = (30, 60, 90), takip: int = 90, buyuk: Decimal = Decimal(250000)) -> dict:
    krediler, uy = oku(yol)
    r = degerlendir(krediler, dilimler, takip, buyuk)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, takip)
    return r


def dilim_coz(s: str) -> tuple[int, ...] | None:
    try:
        d = tuple(int(x) for x in re.split(r"[;,\s]+", s.strip()) if x)
    except ValueError:
        return None
    return d if d and d[0] >= 1 and all(a < b for a, b in zip(d, d[1:])) else None


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Gecikmeye düşen ve erken uyarı veren kredileri tutar ve gün bazında önceliklendirir.")
    p.add_argument("--krediler", type=Path, help="Müşteri No, Müşteri, Kredi No, Ürün, Risk Bakiyesi, Geciken Tutar, Gecikme Gün, Erken Uyarı, Portföy Yöneticisi")
    p.add_argument("--dilimler", default="30;60;90", help="Gecikme dilimi sınırları, gün (varsayılan 30;60;90)")
    p.add_argument("--takip", type=int, default=90, help="Takip eşiği, gün; kalan gün buna göre hesaplanır (varsayılan 90)")
    p.add_argument("--buyuk", default="250000", help="Geciken tutar bu değer ve üstündeyse müşteri en az 'Yüksek' öncelik alır (varsayılan 250.000)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "sube_risk_gecikme_listesi.xlsx")
    a = p.parse_args(argv)
    yol = a.krediler or ORNEK / "krediler.csv"
    dilimler, buyuk = dilim_coz(a.dilimler), para(a.buyuk)
    if dilimler is None or a.takip < 1 or buyuk is None or buyuk <= 0:
        print("[X] --dilimler artan pozitif tam sayılar (örn. 30;60;90); --takip en az 1; --buyuk pozitif olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, dilimler, a.takip, buyuk)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t, s = r["toplam"], r["say"]
    print(f"[OK] {len(r['krediler'])} kredi · risk {t['risk']:,.2f} · geciken {t['geciken']:,.2f}" + (f" · gecikme oranı %{t['oran']}" if t["oran"] is not None else "")
          + f" · Kritik {s['1 – Kritik']}, Yüksek {s['2 – Yüksek']}, Orta {s['3 – Orta']}, Erken uyarı {s['4 – Erken uyarı']}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
