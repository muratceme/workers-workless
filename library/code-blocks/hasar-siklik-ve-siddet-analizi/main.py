"""
Hasar Sıklık ve Şiddet Analizi — Workers / Workless kod bloğu
Sigortacılık › Aktüerya › Aktüerya Uzmanı

Poliçe (maruziyet) ve hasar dökümünden, dönem ve segment bazında:
  - Kazanılmış maruziyet (poliçe-yıl) ve kazanılmış prim: poliçe süresinin döneme düşen gün oranıyla.
  - Hasar adedi, sıklık (adet / maruziyet), şiddet (gerçekleşen / adet), saf prim (gerçekleşen / maruziyet),
    hasar/prim oranı (gerçekleşen / kazanılmış prim). Gerçekleşen = ödenen + muallak.
  - Büyük hasar ayrımı: eşik (tutar veya yüzdelik) üstü hasarlar sınırlanarak "sınırlı şiddet" de verilir.
  - Sıklık için %95 güven aralığı (Poisson yaklaşımı) ve klasik tam güvenilirlik ölçütüne göre (1.082 hasar;
    p = %90, k = %5) güvenilirlik oranı √(n / 1082).
İnternete bağlanmaz.

Kullanım:
    python main.py                                         # örnek kasko portföyüyle dener
    python main.py --policeler policeler.xlsx --hasarlar hasarlar.xlsx --segment "Araç Tipi" --segment Bölge
    python main.py --policeler p.xlsx --hasarlar h.xlsx --donem ceyrek --buyuk-hasar 500000
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import OrderedDict, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
TAM_GUVENILIRLIK = 1082          # klasik (limited fluctuation) ölçüt: p = 0,90, k = 0,05 → (1,645 / 0,05)² ≈ 1.082


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def katla(s) -> str:
    s = kucuk(s)
    for a, b in zip("ıöüşğç", "iousgc"):
        s = s.replace(a, b)
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def sayi(x) -> float:
    if x in (None, ""):
        return 0.0
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip().replace("TL", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1:
        s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        return 0.0


def tl(x: float, b: int = 0) -> str:
    return f"{x:,.{b}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    for b in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(str(x or "").strip()[:10], b).date()
        except ValueError:
            continue
    return None


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def _bul(baslik, *adlar):
    b = [katla(x) for x in baslik]
    return next((b.index(katla(a)) for a in adlar if katla(a) in b), None)


def _al(r, i):
    return r[i] if i is not None and i < len(r) else None


# ----------------------------------------------------------------------------
# Okuma
# ----------------------------------------------------------------------------

def policeler_oku(yol: Path, segmentler: list[str]) -> list[dict]:
    s = tablo_oku(yol)
    b = s[0]
    i_no, i_bas = _bul(b, "poliçe no", "police no", "poliçe"), _bul(b, "başlangıç", "başlangıç tarihi", "vade başlangıcı")
    i_bit, i_prim = _bul(b, "bitiş", "bitiş tarihi", "vade bitişi"), _bul(b, "prim", "brüt prim", "net prim", "tahakkuk eden prim")
    i_ipt = _bul(b, "iptal tarihi", "iptal")
    if None in (i_no, i_bas, i_bit):
        raise SystemExit(f"Poliçe dosyasında Poliçe No, Başlangıç ve Bitiş sütunları gerekli. Başlıklar: {b}")
    i_seg = {sg: _bul(b, sg) for sg in segmentler}
    eksik = [sg for sg, i in i_seg.items() if i is None]
    if eksik:
        raise SystemExit(f"Poliçe dosyasında segment sütunu yok: {', '.join(eksik)}. Başlıklar: {b}")
    sonuc = []
    for r in s[1:]:
        bas, bit = tarih(_al(r, i_bas)), tarih(_al(r, i_bit))
        if not _al(r, i_no) or not bas or not bit or bit <= bas:
            continue
        ipt = tarih(_al(r, i_ipt))
        sonuc.append({"no": str(_al(r, i_no)).strip(), "bas": bas, "bit": min(bit, ipt) if ipt else bit, "sure": (bit - bas).days,
                      "prim": sayi(_al(r, i_prim)), "seg": {sg: str(_al(r, i) or "-").strip() for sg, i in i_seg.items()}})
    return sonuc


def hasarlar_oku(yol: Path) -> list[dict]:
    s = tablo_oku(yol)
    b = s[0]
    i_p, i_t = _bul(b, "poliçe no", "police no", "poliçe"), _bul(b, "hasar tarihi", "kaza tarihi", "olay tarihi")
    i_o, i_m = _bul(b, "ödenen", "ödenen tutar", "ödeme"), _bul(b, "muallak", "muallak tutar", "rezerv")
    i_g = _bul(b, "gerçekleşen", "toplam hasar", "hasar tutarı")
    i_d = _bul(b, "durum")
    if i_p is None or i_t is None or (i_g is None and i_o is None):
        raise SystemExit(f"Hasar dosyasında Poliçe No, Hasar Tarihi ve Ödenen/Muallak (veya Gerçekleşen) gerekli. Başlıklar: {b}")
    sonuc = []
    for n, r in enumerate(s[1:], 1):
        t = tarih(_al(r, i_t))
        if not t or not _al(r, i_p):
            continue
        if katla(_al(r, i_d)) in ("red", "reddedildi", "ret", "iptal"):
            continue                                          # reddedilen dosyalar sıklığa girmez
        tutar = sayi(_al(r, i_g)) if i_g is not None else sayi(_al(r, i_o)) + sayi(_al(r, i_m))
        sonuc.append({"police": str(_al(r, i_p)).strip(), "tarih": t, "tutar": tutar})
    return sonuc


# ----------------------------------------------------------------------------
# Hesap
# ----------------------------------------------------------------------------

def donem_anahtari(d: date, tur: str) -> str:
    return f"{d.year}" if tur == "yil" else f"{d.year}-Ç{(d.month - 1) // 3 + 1}"


def donem_sinirlari(anahtar: str, tur: str) -> tuple[date, date]:
    if tur == "yil":
        y = int(anahtar)
        return date(y, 1, 1), date(y + 1, 1, 1)
    y, c = int(anahtar[:4]), int(anahtar[-1])
    bas = date(y, 3 * (c - 1) + 1, 1)
    son = date(y + (1 if c == 4 else 0), 1 if c == 4 else 3 * c + 1, 1)
    return bas, son


def kazanilmis(policeler: list[dict], tur: str, degerleme: date) -> dict[tuple[str, tuple], list[float]]:
    """(dönem, segment değerleri) → [maruziyet (poliçe-yıl), kazanılmış prim]. Poliçe günleri değerleme tarihine kadar
    dönemlere oransal dağıtılır (değerleme tarihinden sonrası kazanılmamıştır)."""
    sonuc: dict[tuple[str, tuple], list[float]] = defaultdict(lambda: [0.0, 0.0])
    for p in policeler:
        d = p["bas"]
        bitis = min(p["bit"], degerleme + timedelta(days=1))
        while d < bitis:
            anahtar = donem_anahtari(d, tur)
            _, donem_son = donem_sinirlari(anahtar, tur)
            parca_son = min(donem_son, bitis)
            gun = (parca_son - d).days
            x = sonuc[(anahtar, tuple(p["seg"].values()))]
            x[0] += gun / 365.0
            x[1] += p["prim"] * gun / p["sure"]
            d = parca_son
    return sonuc


def esik_hesapla(tutarlar: list[float], buyuk_hasar: str | None) -> float | None:
    if not buyuk_hasar or not tutarlar:
        return None
    m = re.fullmatch(r"p(\d{1,2}(?:\.\d+)?)", buyuk_hasar.strip().lower())
    if m:
        sirali = sorted(tutarlar)
        k = (len(sirali) - 1) * float(m.group(1)) / 100
        alt, ust = math.floor(k), math.ceil(k)
        return sirali[alt] + (sirali[ust] - sirali[alt]) * (k - alt)
    return sayi(buyuk_hasar)


def analiz(policeler, hasarlar, segmentler, tur, buyuk_hasar, degerleme: date) -> dict:
    kz = kazanilmis(policeler, tur, degerleme)
    hasarlar = [h for h in hasarlar if h["tarih"] <= degerleme]
    pmap = {p["no"]: p for p in policeler}
    esik = esik_hesapla([h["tutar"] for h in hasarlar], buyuk_hasar)
    eslesmeyen = 0
    hasar_gr: dict[tuple[str, tuple], list] = defaultdict(list)
    for h in hasarlar:
        p = pmap.get(h["police"])
        if not p:
            eslesmeyen += 1
            continue
        hasar_gr[(donem_anahtari(h["tarih"], tur), tuple(p["seg"].values()))].append(h["tutar"])

    def topla(anahtar_fn):
        sonuc = defaultdict(lambda: {"maruziyet": 0.0, "prim": 0.0, "tutarlar": []})
        for (dn, sg), (e, pr) in kz.items():
            x = sonuc[anahtar_fn(dn, sg)]
            x["maruziyet"] += e
            x["prim"] += pr
        for (dn, sg), ts in hasar_gr.items():
            sonuc[anahtar_fn(dn, sg)]["tutarlar"].extend(ts)
        return {k: olcutler(v, esik) for k, v in sorted(sonuc.items())}

    return {"donem": topla(lambda dn, sg: dn), "segment": {sg: topla(lambda dn, s, i=i: s[i]) for i, sg in enumerate(segmentler)},
            "capraz": topla(lambda dn, sg: (dn,) + sg) if segmentler else {}, "esik": esik, "eslesmeyen": eslesmeyen,
            "toplam": olcutler({"maruziyet": sum(v[0] for v in kz.values()), "prim": sum(v[1] for v in kz.values()),
                                "tutarlar": [t for ts in hasar_gr.values() for t in ts]}, esik)}


def olcutler(x: dict, esik: float | None) -> dict:
    n, e, prim = len(x["tutarlar"]), x["maruziyet"], x["prim"]
    toplam = sum(x["tutarlar"])
    sinirli = sum(min(t, esik) for t in x["tutarlar"]) if esik else toplam
    siklik = n / e if e else None
    sd = math.sqrt(n) / e if e else None
    return {"maruziyet": e, "prim": prim, "adet": n, "toplam": toplam, "siklik": siklik,
            "siklik_alt": max(0.0, siklik - 1.96 * sd) if siklik is not None else None,
            "siklik_ust": siklik + 1.96 * sd if siklik is not None else None,
            "siddet": toplam / n if n else None, "siddet_sinirli": sinirli / n if n else None,
            "saf_prim": toplam / e if e else None, "hp_orani": toplam / prim if prim else None,
            "buyuk": sum(1 for t in x["tutarlar"] if esik and t > esik),
            "guvenilirlik": min(1.0, math.sqrt(n / TAM_GUVENILIRLIK))}


def calistir(police_yolu: Path, hasar_yolu: Path, cikti: Path, segmentler: list[str] | None = None, tur: str = "yil",
             buyuk_hasar: str | None = "p95", degerleme: date | None = None) -> dict:
    segmentler = segmentler or []
    policeler = policeler_oku(police_yolu, segmentler)
    hasarlar = hasarlar_oku(hasar_yolu)
    degerleme = degerleme or date.today()
    a = analiz(policeler, hasarlar, segmentler, tur, buyuk_hasar, degerleme)
    a["degerleme"] = degerleme
    _rapor(a, segmentler, tur, cikti)
    return a


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
SARI = PatternFill("solid", fgColor="FFF4CE")
SUTUN = ["Maruziyet (poliçe-yıl)", "Kazanılmış Prim", "Hasar Adedi", "Gerçekleşen Hasar", "Sıklık", "Sıklık %95 Alt", "Sıklık %95 Üst",
         "Şiddet", "Sınırlı Şiddet", "Saf Prim", "Hasar/Prim", "Büyük Hasar", "Güvenilirlik"]
ALAN = ["maruziyet", "prim", "adet", "toplam", "siklik", "siklik_alt", "siklik_ust", "siddet", "siddet_sinirli", "saf_prim", "hp_orani",
        "buyuk", "guvenilirlik"]
BICIM = ["#,##0.0", "#,##0", "0", "#,##0", "0.0000", "0.0000", "0.0000", "#,##0", "#,##0", "#,##0.00", "0.0%", "0", "0%"]


def _baslik(ws, satir=1):
    for c in ws[satir]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI


def _tablo(ws, baslik: list[str], veri: dict, ilk_sutun_genislik: int = 14):
    ws.append(baslik + SUTUN)
    _baslik(ws, ws.max_row)
    for k, v in veri.items():
        anahtar = list(k) if isinstance(k, tuple) else [k]
        ws.append(anahtar + [v[a] for a in ALAN])
        for j, bc in enumerate(BICIM, len(anahtar) + 1):
            ws.cell(ws.max_row, j).number_format = bc
        if v["guvenilirlik"] < 0.5:
            ws.cell(ws.max_row, len(anahtar) + len(ALAN)).fill = SARI
    for j in range(1, len(baslik) + 1):
        ws.column_dimensions[get_column_letter(j)].width = ilk_sutun_genislik
    for j in range(len(baslik) + 1, len(baslik) + len(ALAN) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 13


def _rapor(a, segmentler, tur, cikti):
    wb = Workbook()
    o = wb.active
    o.title = "Dönemler"
    o.append(["Hasar sıklık ve şiddet analizi — dönem bazında"])
    o["A1"].font = Font(bold=True, size=12)
    _tablo(o, ["Dönem"], a["donem"])
    t = a["toplam"]
    o.append(["TOPLAM"] + [t[x] for x in ALAN])
    for j, bc in enumerate(BICIM, 2):
        o.cell(o.max_row, j).number_format = bc
    for c in o[o.max_row]:
        c.font = Font(bold=True)
    n = len(a["donem"])
    if n >= 2:
        g = BarChart()
        g.title, g.height, g.width = "Sıklık (bar) ve şiddet (çizgi)", 8, 22
        g.add_data(Reference(o, min_col=6, min_row=2, max_row=2 + n), titles_from_data=True)
        g.set_categories(Reference(o, min_col=1, min_row=3, max_row=2 + n))
        c = LineChart()
        c.add_data(Reference(o, min_col=9, min_row=2, max_row=2 + n), titles_from_data=True)
        c.y_axis.axId = 200
        c.y_axis.crosses = "max"
        g += c
        o.add_chart(g, f"B{n + 6}")
    for sg in segmentler:
        ws = wb.create_sheet(sg[:28])
        _tablo(ws, [sg], a["segment"][sg], 18)
    if segmentler:
        ws = wb.create_sheet("Dönem × Segment")
        _tablo(ws, ["Dönem"] + segmentler, a["capraz"])
    b = wb.create_sheet("Bilgi")
    esik = a["esik"]
    for s in [["Maruziyet", "poliçe süresinin döneme düşen günü / 365 (poliçe-yıl); iptal tarihi varsa orada biter"],
              ["Kazanılmış prim", "poliçe primi × döneme düşen gün / poliçe süresi (gün)"],
              ["Gerçekleşen", "ödenen + muallak (reddedilen dosyalar hariç); IBNR dahil değildir — son dönemler eksik gelişmiştir"],
              ["Sıklık", "hasar adedi / maruziyet; %95 aralık: sıklık ± 1,96 × √adet / maruziyet (Poisson yaklaşımı)"],
              ["Şiddet", "gerçekleşen / adet; sınırlı şiddet büyük hasar eşiğinde sınırlanmış tutarlarla"],
              ["Büyük hasar eşiği", tl(esik) if esik else "uygulanmadı"],
              ["Güvenilirlik", f"√(adet / {TAM_GUVENILIRLIK}) — klasik tam güvenilirlik ölçütü (p = %90, k = %5); %50'nin altı sarı"],
              ["Eşleşmeyen hasar", f"{a['eslesmeyen']} hasarın poliçesi poliçe dosyasında bulunamadı"],
              ["Not", "Son dönemlerin hasarları henüz tam gelişmediği için sıklık ve şiddet düşük görünebilir; fiyatlamada IBNR ve "
                      "gelişim faktörleriyle düzeltin (bkz. IBNR Rezerv Tahmini paketi)."]]:
        b.append(s)
    b.column_dimensions["A"].width = 20
    b.column_dimensions["B"].width = 130
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ornek = BURASI / "ornek_veri"
    ap = argparse.ArgumentParser(description="Poliçe ve hasar dökümünden dönem ve segment bazında sıklık, şiddet ve hasar/prim analizi yapar.")
    ap.add_argument("--policeler", type=Path, default=ornek / "policeler.csv",
                    help="Poliçeler: Poliçe No, Başlangıç, Bitiş, Prim [, İptal Tarihi, segment sütunları]")
    ap.add_argument("--hasarlar", type=Path, default=ornek / "hasarlar.csv",
                    help="Hasarlar: Hasar No, Poliçe No, Hasar Tarihi, Ödenen, Muallak [, Durum] — veya Gerçekleşen")
    ap.add_argument("--segment", action="append", help="Kırılım sütunu (birden çok verilebilir), ör. --segment 'Araç Tipi'")
    ap.add_argument("--donem", choices=["yil", "ceyrek"], default="yil", help="Dönem (varsayılan yıl)")
    ap.add_argument("--buyuk-hasar", default="p95", help="Büyük hasar eşiği: tutar (ör. 500000) veya yüzdelik (ör. p95, p99); 'yok' = uygulanmaz")
    ap.add_argument("--degerleme", help="Değerleme tarihi GG.AA.YYYY (varsayılan bugün): maruziyet ve hasarlar bu tarihe kadar")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "siklik_siddet.xlsx")
    a = ap.parse_args(argv)
    if a.policeler == ornek / "policeler.csv" and not a.segment:
        a.segment = ["Araç Tipi", "Bölge"]
    bh = None if katla(a.buyuk_hasar) in ("yok", "") else a.buyuk_hasar
    degerleme = tarih(a.degerleme) if a.degerleme else (date(2026, 9, 30) if a.policeler == ornek / "policeler.csv" else None)
    s = calistir(a.policeler, a.hasarlar, a.cikti, a.segment, a.donem, bh, degerleme)
    t = s["toplam"]
    print(f"[OK] maruziyet {tl(t['maruziyet'], 1)} poliçe-yıl · {t['adet']} hasar · sıklık {tl(t['siklik'] or 0, 4)} · "
          f"şiddet {tl(t['siddet'] or 0)} · hasar/prim %{tl((t['hp_orani'] or 0) * 100, 1)}")
    for d, v in s["donem"].items():
        print(f"    {d}: sıklık {tl(v['siklik'] or 0, 4)} · şiddet {tl(v['siddet'] or 0)} · H/P %{tl((v['hp_orani'] or 0) * 100, 1)}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
