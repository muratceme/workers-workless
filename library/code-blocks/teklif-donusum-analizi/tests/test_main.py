import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
GUN = date(2026, 9, 30)


def t(no, ay, gun, musteri, hat, durum, fiyat, rakip=None, neden=""):
    d = date(2026, ay, gun)
    return {"no": no, "tarih": d, "musteri": musteri, "hat": hat, "mod": "FCL", "fiyat": D(fiyat), "para": "USD", "durum": durum, "rakip": None if rakip is None else D(rakip), "neden": neden,
            "temsilci": "S", "ay": f"{d:%Y-%m}"}


def veri():
    return [t("T1", 9, 1, "C1", "A", "Kazanıldı", 1000, 1050), t("T2", 9, 2, "C1", "A", "Kaybedildi", 1000, 900, "Fiyat"), t("T3", 9, 3, "C2", "a", "Kaybedildi", 1200, 1100, "fiyat"),
            t("T4", 9, 3, "C2", "B", "Kazanıldı", 500), t("T5", 8, 5, "C3", "B", "Beklemede", 800), t("T6", 9, 20, "C3", "B", "Beklemede", 700), t("T7", 8, 10, "C1", "A", "İptal", 900)]


def k(s):
    return (s["ad"], s["teklif"], s["kazanilan"], s["kaybedilen"], s["beklemede"], s["iptal"], s["oran"], s["tutar"])


class AnalizTesti(unittest.TestCase):
    def test_oranlar(self):
        r = main.analiz(veri(), GUN)
        self.assertEqual(k({"ad": "-", **r["toplam"]}), ("-", 7, 2, 2, 2, 1, D("50.00"), {"USD": D(1500)}))
        self.assertEqual([k(s) for s in r["kirilim"]["Hat"]], [("A", 4, 1, 2, 0, 1, D("33.33"), {"USD": D(1000)}), ("B", 3, 1, 0, 2, 0, D("100.00"), {"USD": D(500)})])
        self.assertEqual([(s["ad"], s["oran"]) for s in r["kirilim"]["Ay"]], [("2026-08", None), ("2026-09", D("50.00"))])
        self.assertEqual([(b["bant"], b["teklif"], b["oran"]) for b in r["bantlar"]], [("≤ 0 (rakipten ucuz ya da eşit)", 1, D("100.00")), ("0–5", 0, None), ("5–10", 1, D("0.00")), ("10 üstü", 1, D("0.00"))])
        self.assertEqual(([(f["no"], f["fark"]) for f in r["farkli"]], r["kayip_ort_fark"], r["nedenler"]), ([("T1", D("-4.76")), ("T2", D("11.11")), ("T3", D("9.09"))], D("10.10"), [{"neden": "Fiyat", "adet": 2, "pay": D("100.00")}]))
        self.assertEqual(([(a["no"], a["gun"]) for a in r["acik"]], [(u["tur"], u["kim"]) for u in r["uyarilar"]]), ([("T5", 56)], [("Sonuç bekleyen eski teklif", "1 teklif")]))

    def test_esikler(self):
        r = main.analiz(veri(), GUN, en_az=3, esik=D(50), bekleme=60)
        self.assertEqual([(u["tur"], u["kim"]) for u in r["uyarilar"]], [("Düşük kazanma oranı", "Hat: A")])
        self.assertEqual([main._durum(x) for x in ("", "Kazanıldı", "KAYBEDİLDİ", "Reddedildi", "İptal edildi", "Açık", "Değerlendirmede", "x")],
                         ["Beklemede", "Kazanıldı", "Kaybedildi", "Kaybedildi", "İptal", "Beklemede", "Beklemede", None])
        with self.assertRaises(ValueError):
            main.analiz([], GUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "navlun_teklifleri.csv", cikti, GUN)
            self.assertEqual((k({"ad": "-", **r["toplam"]}), r["kayip_ort_fark"]), (("-", 62, 15, 36, 6, 5, D("29.41"), {"EUR": D("17089.25"), "USD": D("12572.41")}), D("6.86")))
            self.assertEqual([(b["teklif"], b["oran"]) for b in r["bantlar"]], [(11, D("63.64")), (7, D("0.00")), (9, D("0.00")), (9, D("0.00"))])
            self.assertEqual([(n["neden"], n["adet"]) for n in r["nedenler"]], [("Fiyat", 22), ("Müşteri vazgeçti", 6), ("Transit süre", 6), ("Ekipman / araç yok", 2)])
            self.assertEqual([(u["tur"], u["kim"]) for u in r["uyarilar"]], [("Düşük kazanma oranı", "Hat: Gebze – Paris"), ("Düşük kazanma oranı", "Hat: İzmir – Rotterdam")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Kırılımlar", "Fiyat Farkı", "Kayıp Nedenleri", "Bekleyen Teklifler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "t.csv"
            y.write_text("Teklif Numarası;Teklif Tarihi;Firma;POL;POD;Navlun;Döviz;Sonuç;Kazanan Fiyat\nQ1;01.09.2026;F;Mersin;Hamburg;1.250,00;usd;Kaybedildi;1.100\nQ1;01.09.2026;F;;;5;;;\n"
                         "Q2;x;F;;;5;;;\nQ3;01.09.2026;F;;;0;;;\nQ4;01.09.2026;F;;;5;;belki;\nQ5;01.09.2026;F;;;5;;;y\nQ6;01.09.2026;;İzmir;;5;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["hat"], x["para"], x["durum"], x["rakip"], x["musteri"]) for x in l], [u["tur"] for u in uy]),
                             ([("Q1", "Mersin – Hamburg", "USD", "Kaybedildi", D(1100), "F"), ("Q6", "İzmir", "TRY", "Beklemede", None, "Belirtilmemiş")],
                              ["Teklif iki kez yazılmış", "Satır okunamadı", "Satır okunamadı", "Satır okunamadı", "Satır okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--teklifler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "30.09.2026", "--min", "1", "--esik", "50", "--bekleme", "7"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--min", "0"], ["--esik", "101"], ["--bekleme", "-1"])], [2] * 4)
            self.assertEqual(main.main(["--teklifler", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Teklif No;Firma\nQ1;F\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
