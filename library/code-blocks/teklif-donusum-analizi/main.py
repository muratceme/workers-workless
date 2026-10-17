"""
Teklif Dönüşüm Analizi — Gerisi Bende kod bloğu
Satış ve Fiyatlandırma › Lojistik Satış Uzmanı
Verilen navlun tekliflerinin kazanılma oranını hat, müşteri, ay, taşıma modu ve temsilci bazında; kaybedilenleri fiyat farkı ve nedene göre analiz eder.
  Kazanma oranı %   = kazanılan teklif ÷ (kazanılan + kaybedilen) × 100   (beklemede ve iptal teklifler oranın dışında)
  Fiyat farkı %     = (teklif fiyatımız − rakip fiyatı) ÷ rakip fiyatı × 100   (yalnız aynı para biriminde rakip fiyatı yazılmışsa)
  Fiyat farkı bandı = ≤ 0 · 0–5 · 5–10 · 10 üstü; her bantta kazanma oranı
  Uyarı             = en az --min (5) sonuçlanmış teklifi olan hat / müşteride kazanma oranı --esik (%20) altında
                      · --bekleme (14) günden uzun süredir sonuç bekleyen teklif
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"no": ("teklif no", "teklif numarasi", "teklif", "no"), "tarih": ("tarih", "teklif tarihi"), "musteri": ("musteri", "musteri adi", "firma", "cari"),
         "cikis": ("cikis", "yukleme yeri", "cikis noktasi", "pol", "cikis limani"), "varis": ("varis", "teslim yeri", "varis noktasi", "pod", "varis limani"), "hat": ("hat", "guzergah", "rota"),
         "mod": ("mod", "tasima modu", "tasima turu", "servis"), "fiyat": ("teklif fiyati", "fiyat", "navlun", "teklif tutari", "tutar"), "para": ("para birimi", "doviz", "pb", "para"),
         "durum": ("durum", "sonuc", "teklif durumu"), "rakip": ("rakip fiyat", "rakip fiyati", "kazanan fiyat"), "neden": ("kayip nedeni", "kaybetme nedeni", "neden", "aciklama"),
         "temsilci": ("satis temsilcisi", "temsilci", "satisci", "sorumlu")}
DURUM = (("Kazanıldı", ("kazan", "onay", "won", "kabul")), ("Kaybedildi", ("kayb", "kaybet", "red", "lost", "olumsuz")), ("İptal", ("iptal", "vazgec")),
         ("Beklemede", ("bekle", "acik", "degerlendir", "pending", "teklif verildi")))
BANT = (("≤ 0 (rakipten ucuz ya da eşit)", Decimal(0)), ("0–5", Decimal(5)), ("5–10", Decimal(10)), ("10 üstü", None))
KIRILIM = (("hat", "Hat"), ("musteri", "Müşteri"), ("ay", "Ay"), ("mod", "Taşıma Modu"), ("temsilci", "Temsilci"))


def _durum(x) -> str | None:
    k = katla(x)
    if not k:
        return "Beklemede"
    return next((d for d, anahtar in DURUM if any(k.startswith(a) or f" {a}" in f" {k}" for a in anahtar)), None)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("tarih", "fiyat")):
        no = metin(r["no"])
        musteri = metin(r["musteri"])
        if not no and not musteri:
            continue
        t, fiyat, durum = tarih(r["tarih"]), para(r["fiyat"]), _durum(r["durum"])
        rakip = para(r["rakip"]) if metin(r["rakip"]) else None
        kim = no or f"Satır {r['_satir']}"
        if t is None or fiyat is None or fiyat <= 0 or durum is None or (metin(r["rakip"]) and (rakip is None or rakip <= 0)):
            uy.append(uyari("Yüksek", "Satır okunamadı", kim, f"Satır {r['_satir']}: tarih, fiyat, rakip fiyatı ya da durum ('{metin(r['durum'])}') okunamadı. Teklif alınmadı."))
            continue
        if no and katla(no) in gorulen:
            uy.append(uyari("Orta", "Teklif iki kez yazılmış", kim, f"Satır {r['_satir']}: ikinci satır alınmadı; revize teklifse yeni numara verin."))
            continue
        if no:
            gorulen.add(katla(no))
        cikis, varis = metin(r["cikis"]), metin(r["varis"])
        hat = f"{cikis} – {varis}" if cikis and varis else metin(r["hat"]) or cikis or varis or "Belirtilmemiş"
        l.append({"no": no, "tarih": t, "musteri": musteri or "Belirtilmemiş", "hat": hat, "mod": metin(r["mod"]) or "Belirtilmemiş", "fiyat": fiyat, "para": metin(r["para"]).upper() or "TRY",
                  "durum": durum, "rakip": rakip, "neden": metin(r["neden"]), "temsilci": metin(r["temsilci"]) or "Belirtilmemiş", "ay": f"{t:%Y-%m}"})
    return l, uy


def _oran(l: list[dict]) -> dict:
    k, y = sum(1 for t in l if t["durum"] == "Kazanıldı"), sum(1 for t in l if t["durum"] == "Kaybedildi")
    tutar = defaultdict(lambda: Decimal(0))
    for t in l:
        if t["durum"] == "Kazanıldı":
            tutar[t["para"]] += t["fiyat"]
    return {"teklif": len(l), "kazanilan": k, "kaybedilen": y, "beklemede": sum(1 for t in l if t["durum"] == "Beklemede"), "iptal": sum(1 for t in l if t["durum"] == "İptal"),
            "oran": yuvarla(Decimal(k) * 100 / (k + y)) if k + y else None, "tutar": dict(sorted(tutar.items()))}


def analiz(teklifler: list[dict], gun: date, en_az: int = 5, esik: Decimal = Decimal(20), bekleme: int = 14) -> dict:
    if not teklifler:
        raise ValueError("Teklif kaydı yok.")
    uy, kirilim = [], {}
    for alan, ad in KIRILIM:
        g = {}
        for t in teklifler:
            g.setdefault(katla(t[alan]), []).append(t)
        satirlar = sorted(({"ad": l[0][alan], **_oran(l)} for l in g.values()), key=lambda s: (s["ad"] if alan == "ay" else "", -s["teklif"], katla(s["ad"])))
        kirilim[ad] = satirlar
        if alan in ("hat", "musteri"):
            for s in satirlar:
                if s["kazanilan"] + s["kaybedilen"] >= en_az and s["oran"] < esik:
                    uy.append(uyari("Orta", "Düşük kazanma oranı", f"{ad}: {s['ad']}", f"{s['kazanilan'] + s['kaybedilen']} sonuçlanan tekliften {s['kazanilan']} kazanıldı (%{s['oran']}; eşik %{esik})."))
    farkli = []
    for t in teklifler:
        if t["rakip"] is not None and t["durum"] in ("Kazanıldı", "Kaybedildi"):
            fark = yuvarla((t["fiyat"] - t["rakip"]) * 100 / t["rakip"])
            farkli.append({**t, "fark": fark, "bant": next(ad for ad, ust in BANT if ust is None or fark <= ust)})
    bantlar = [{"bant": ad, **_oran([t for t in farkli if t["bant"] == ad])} for ad, _ in BANT]
    kayip = [t for t in teklifler if t["durum"] == "Kaybedildi"]
    nedenler = {}
    for t in kayip:
        nedenler.setdefault(katla(t["neden"]) or "-", []).append(t)
    neden = sorted(({"neden": l[0]["neden"] or "Belirtilmemiş", "adet": len(l), "pay": yuvarla(Decimal(len(l)) * 100 / len(kayip))} for l in nedenler.values()), key=lambda n: (-n["adet"], katla(n["neden"])))
    kayip_fark = [t["fark"] for t in farkli if t["durum"] == "Kaybedildi"]
    acik = []
    for t in teklifler:
        if t["durum"] == "Beklemede":
            g = (gun - t["tarih"]).days
            if g > bekleme:
                acik.append({**t, "gun": g})
    acik.sort(key=lambda t: (-t["gun"], t["no"]))
    if acik:
        uy.append(uyari("Orta", "Sonuç bekleyen eski teklif", f"{len(acik)} teklif", f"{bekleme} günden uzun süredir sonucu girilmemiş teklifler var; müşteriyle takip edin ya da durumunu güncelleyin."))
    return {"toplam": _oran(teklifler), "kirilim": kirilim, "bantlar": bantlar, "farkli": farkli, "nedenler": neden, "kayip_ort_fark": yuvarla(sum(kayip_fark) / len(kayip_fark)) if kayip_fark else None,
            "acik": acik, "uyarilar": uy}


def _tutar(d: dict) -> str:
    return " · ".join(f"{v:,.2f} {k}" for k, v in d.items())


def rapor_yaz(cikti: Path, r: dict, gun: date, esik: Decimal, bekleme: int) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    t = r["toplam"]
    _baslik(ws, ["Gösterge", "Değer"], (40, 30))
    for ad, v in (("Teklif sayısı", t["teklif"]), ("Kazanılan", t["kazanilan"]), ("Kaybedilen", t["kaybedilen"]), ("Beklemede", t["beklemede"]), ("İptal", t["iptal"]),
                  ("Kazanma oranı % (kazanılan ÷ sonuçlanan)", t["oran"]), ("Kazanılan teklif tutarı", _tutar(t["tutar"])), ("Kaybedilenlerde ortalama fiyat farkı %", r["kayip_ort_fark"]),
                  (f"{bekleme} günden eski bekleyen teklif", len(r["acik"])), ("Rapor tarihi", gun)):
        _satir(ws, [ad, v], tarih_s=(2,) if isinstance(v, date) else ())
    ws.append([])
    ws.append(["Fiyat farkı = (bizim − rakip) ÷ rakip × 100; yalnız rakip fiyatı yazılmış sonuçlanan teklifler."])
    ws = wb.create_sheet("Kırılımlar")
    _baslik(ws, ["Kırılım", "Ad", "Teklif", "Kazanılan", "Kaybedilen", "Beklemede", "İptal", "Kazanma Oranı %", "Kazanılan Tutar"], (13, 30, 8, 9, 10, 10, 7, 10, 26), "C2")
    for ad, satirlar in r["kirilim"].items():
        for s in satirlar:
            _satir(ws, [ad, s["ad"], s["teklif"], s["kazanilan"], s["kaybedilen"], s["beklemede"], s["iptal"], s["oran"], _tutar(s["tutar"])], para_s=(8,),
                   dolgu={8: None if s["oran"] is None else KIRMIZI if s["oran"] < esik else YESIL if s["oran"] >= 50 else None})
    ws = wb.create_sheet("Fiyat Farkı")
    _baslik(ws, ["Fiyat Farkı Bandı %", "Teklif", "Kazanılan", "Kaybedilen", "Kazanma Oranı %"], (30, 8, 9, 10, 12))
    for b in r["bantlar"]:
        _satir(ws, [b["bant"], b["teklif"], b["kazanilan"], b["kaybedilen"], b["oran"]], para_s=(5,))
    ws.append([])
    _baslik(ws, ["Teklif No", "Müşteri", "Hat", "Durum", "Para", "Bizim Fiyat", "Rakip Fiyat", "Fark %"], dondur=None)
    for f in sorted(r["farkli"], key=lambda f: (-f["fark"], f["no"])):
        _satir(ws, [f["no"], f["musteri"], f["hat"], f["durum"], f["para"], f["fiyat"], f["rakip"], f["fark"]], para_s=(6, 7, 8), dolgu={4: YESIL if f["durum"] == "Kazanıldı" else KIRMIZI})
    ws = wb.create_sheet("Kayıp Nedenleri")
    _baslik(ws, ["Neden", "Teklif", "Pay %"], (36, 8, 8))
    for n in r["nedenler"]:
        _satir(ws, [n["neden"], n["adet"], n["pay"]], para_s=(3,))
    ws = wb.create_sheet("Bekleyen Teklifler")
    _baslik(ws, ["Teklif No", "Tarih", "Geçen Gün", "Müşteri", "Hat", "Fiyat", "Para", "Temsilci"], (14, 11, 9, 24, 28, 12, 6, 16))
    for a in r["acik"]:
        _satir(ws, [a["no"], a["tarih"], a["gun"], a["musteri"], a["hat"], a["fiyat"], a["para"], a["temsilci"]], para_s=(6,), tarih_s=(2,))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, gun: date, en_az: int = 5, esik: Decimal = Decimal(20), bekleme: int = 14) -> dict:
    teklifler, uy = oku(yol)
    r = analiz(teklifler, gun, en_az, esik, bekleme)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, gun, esik, bekleme)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Verilen navlun tekliflerinin kazanılma oranını hat, müşteri ve fiyat farkı bazında analiz eder.")
    p.add_argument("--teklifler", type=Path, help="Teklif No, Tarih, Müşteri, Çıkış, Varış (ya da Hat), Mod, Teklif Fiyatı, Para Birimi, Durum, Rakip Fiyat, Kayıp Nedeni, Satış Temsilcisi")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY; bekleyen tekliflerin yaşı için (varsayılan bugün; örnek veride 30.09.2026)")
    p.add_argument("--min", type=int, default=5, help="Kazanma oranı uyarısı için hat / müşteride en az sonuçlanmış teklif (varsayılan 5)")
    p.add_argument("--esik", default="20", help="Kazanma oranı bu yüzdenin altındaysa uyarılır (varsayılan 20)")
    p.add_argument("--bekleme", type=int, default=14, help="Bu kadar günden uzun süredir sonuç bekleyen teklifler listelenir (varsayılan 14)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "teklif_donusum_analizi.xlsx")
    a = p.parse_args(argv)
    yol = a.teklifler or ORNEK / "navlun_teklifleri.csv"
    gun = tarih(a.tarih) if a.tarih else (date(2026, 9, 30) if a.teklifler is None else date.today())
    esik = para(a.esik)
    if gun is None or a.min < 1 or esik is None or not 0 <= esik <= 100 or a.bekleme < 0:
        print("[X] --tarih GG.AA.YYYY, --min en az 1, --esik 0-100 arası, --bekleme sıfır ya da üstü olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, gun, a.min, esik, a.bekleme)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {t['teklif']} teklif · kazanılan {t['kazanilan']} · kaybedilen {t['kaybedilen']} · beklemede {t['beklemede']} · kazanma oranı %{t['oran']}"
          + (f" · kaybedilenlerde ortalama fiyat farkı %{r['kayip_ort_fark']}" if r["kayip_ort_fark"] is not None else ""))
    for b in r["bantlar"]:
        if b["teklif"]:
            print(f"     Fiyat farkı {b['bant']}: {b['teklif']} teklif · kazanma %{b['oran']}")
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
