"""
Teklif Takip Listesi — Workers / Workless kod bloğu
Satış › Satış Destek Uzmanı

Verilen tekliflerin durumunu, bekleme süresini ve dönüşüm oranını izler.
  Bekleme süresi (gün)    = durum tarihi − teklif tarihi   (açık teklifler)
  Takipsiz geçen gün      = durum tarihi − son takip tarihi   (son takip yoksa teklif tarihi)
  Açık teklifin durumu    = Geçerliliği doldu → Geçerlilik yaklaşıyor (5 gün) → Takip gerekli (7 gün) → Bekliyor
  Dönüşüm oranı % (adet)  = kazanılan ÷ (kazanılan + kaybedilen) × 100;   tutar bazında da aynı formül
  Sonuçlanma süresi (gün) = sonuç tarihi − teklif tarihi   (kazanılan ve kaybedilen teklifler)
  Yaş dilimleri           = açık tekliflerin bekleme süresine göre 0–7, 8–15, 16–30, 31+ gün
İptal edilen ve hâlâ açık olan teklifler dönüşüm oranına girmez. Eşikler komut satırından değiştirilebilir.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
SUTUN = {"no": ("teklif no", "teklif numarasi", "no"), "tarih": ("teklif tarihi", "tarih"), "musteri": ("musteri", "firma", "cari"), "tutar": ("tutar", "teklif tutari", "toplam"), "durum": ("durum",),
         "gecerlilik": ("gecerlilik tarihi", "gecerlilik", "son gecerlilik", "opsiyon tarihi"), "takip": ("son takip tarihi", "son takip", "son gorusme"), "sonuc": ("sonuc tarihi", "kapanis tarihi"),
         "temsilci": ("temsilci", "satis temsilcisi", "sorumlu"), "neden": ("kayip nedeni", "neden", "ret nedeni")}
DURUM = {"": "Açık", "acik": "Açık", "bekliyor": "Açık", "beklemede": "Açık", "gonderildi": "Açık", "kazanildi": "Kazanıldı", "onaylandi": "Kazanıldı", "kabul": "Kazanıldı", "siparise donustu": "Kazanıldı", "kaybedildi": "Kaybedildi",
         "reddedildi": "Kaybedildi", "red": "Kaybedildi", "iptal": "İptal", "gecersiz": "İptal"}
ONCELIK = {"Geçerliliği doldu": 0, "Geçerlilik yaklaşıyor": 1, "Takip gerekli": 2, "Bekliyor": 3}
ACIK_RENK = {"Geçerliliği doldu": KIRMIZI, "Geçerlilik yaklaşıyor": SARI, "Takip gerekli": SARI, "Bekliyor": None}
Z, YUZ = Decimal(0), Decimal(100)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("no", "tarih", "tutar")):
        no = metin(r["no"])
        if not no:
            continue
        t, tutar, durum = tarih(r["tarih"]), para(r["tutar"]), DURUM.get(katla(r["durum"]))
        if t is None or tutar is None or tutar < 0 or durum is None:
            uy.append(uyari("Yüksek", "Teklif okunamadı", no, f"Satır {r['_satir']}: tarih '{metin(r['tarih'])}', tutar '{metin(r['tutar'])}', durum '{metin(r['durum'])}'. Durum Açık, Kazanıldı, Kaybedildi veya İptal olmalı. Satır alınmadı."))
            continue
        if katla(no) in gorulen:
            uy.append(uyari("Orta", "Mükerrer teklif numarası", no, f"Aynı numara birden çok satırda (satır {r['_satir']}). İkisi de sayıldı; revizyonsa eskisini İptal yapın."))
        gorulen.add(katla(no))
        l.append({"no": no, "tarih": t, "musteri": metin(r["musteri"]), "tutar": tutar, "durum": durum, "gecerlilik": tarih(r["gecerlilik"]), "takip": tarih(r["takip"]), "sonuc": tarih(r["sonuc"]),
                  "temsilci": metin(r["temsilci"]) or "(belirtilmemiş)", "neden": metin(r["neden"])})
    return l, uy


def _oran(pay, payda) -> Decimal | None:
    return yuvarla(Decimal(pay) * YUZ / Decimal(payda)) if payda else None


def _ozet(liste: list[dict]) -> dict:
    grup = {d: [t for t in liste if t["durum"] == d] for d in ("Açık", "Kazanıldı", "Kaybedildi", "İptal")}
    tutar = {d: sum((t["tutar"] for t in g), Z) for d, g in grup.items()}
    sureler = [t["sure"] for t in grup["Kazanıldı"] + grup["Kaybedildi"] if t["sure"] is not None]
    return {"adet": len(liste), "acik": len(grup["Açık"]), "acik_tutar": tutar["Açık"], "kazanilan": len(grup["Kazanıldı"]), "kazanilan_tutar": tutar["Kazanıldı"], "kaybedilen": len(grup["Kaybedildi"]),
            "kaybedilen_tutar": tutar["Kaybedildi"], "iptal": len(grup["İptal"]), "donusum": _oran(len(grup["Kazanıldı"]), len(grup["Kazanıldı"]) + len(grup["Kaybedildi"])),
            "donusum_tutar": _oran(tutar["Kazanıldı"], tutar["Kazanıldı"] + tutar["Kaybedildi"]), "ort_sure": yuvarla(Decimal(sum(sureler)) / len(sureler)) if sureler else None}


def degerlendir(teklifler: list[dict], bugun: date, takip_gun: int = 7, gecerlilik_gun: int = 5, dilimler: tuple[int, ...] = (7, 15, 30)) -> dict:
    if not teklifler:
        raise ValueError("Değerlendirilecek teklif yok.")
    uy = []
    etiketler = [f"{a}–{b}" for a, b in zip((0,) + tuple(d + 1 for d in dilimler), dilimler)] + [f"{dilimler[-1] + 1}+"]
    for t in teklifler:
        kim = f"{t['no']} · {t['musteri']}".strip(" ·")
        t["bekleme"] = t["takipsiz"] = t["sure"] = t["acik_durum"] = t["dilim"] = None
        if t["tarih"] > bugun:
            uy.append(uyari("Orta", "Teklif tarihi ileride", kim, f"Teklif tarihi {t['tarih']:%d.%m.%Y}, durum tarihi {bugun:%d.%m.%Y}."))
        if t["durum"] == "Açık":
            t["bekleme"] = max((bugun - t["tarih"]).days, 0)
            t["takipsiz"] = max((bugun - max(t["takip"] or t["tarih"], t["tarih"])).days, 0)
            t["dilim"] = etiketler[next((i for i, d in enumerate(dilimler) if t["bekleme"] <= d), len(dilimler))]
            kalan = (t["gecerlilik"] - bugun).days if t["gecerlilik"] else None
            if kalan is not None and kalan < 0:
                t["acik_durum"] = "Geçerliliği doldu"
                uy.append(uyari("Yüksek", "Geçerliliği dolmuş açık teklif", kim, f"{t['tutar']:,.2f} tutarındaki teklifin geçerliliği {t['gecerlilik']:%d.%m.%Y} tarihinde doldu ({-kalan} gün önce). Yenileyin ya da kapatın."))
            elif kalan is not None and kalan <= gecerlilik_gun:
                t["acik_durum"] = "Geçerlilik yaklaşıyor"
                uy.append(uyari("Orta", "Geçerlilik süresi doluyor", kim, f"{t['tutar']:,.2f} tutarındaki teklifin geçerliliği {t['gecerlilik']:%d.%m.%Y} tarihinde doluyor ({kalan} gün kaldı)."))
            elif t["takipsiz"] >= takip_gun:
                t["acik_durum"] = "Takip gerekli"
                uy.append(uyari("Orta", "Takip gerekli", kim, f"{t['tutar']:,.2f} tutarındaki teklif {t['takipsiz']} gündür aranmadı" + (f" (son takip {t['takip']:%d.%m.%Y})." if t["takip"] else " (hiç takip kaydı yok).")))
            else:
                t["acik_durum"] = "Bekliyor"
        elif t["durum"] in ("Kazanıldı", "Kaybedildi"):
            if t["sonuc"] and t["sonuc"] >= t["tarih"]:
                t["sure"] = (t["sonuc"] - t["tarih"]).days
            else:
                uy.append(uyari("Bilgi", "Sonuç tarihi yok", kim, "Sonuç tarihi boş ya da teklif tarihinden önce; sonuçlanma süresine katılmadı."))
            if t["durum"] == "Kaybedildi" and not t["neden"]:
                uy.append(uyari("Bilgi", "Kayıp nedeni yazılmamış", kim, "Kaybedilen teklifin nedeni boş; neden dağılımında '(belirtilmemiş)' olarak yer aldı."))
    acik = sorted((t for t in teklifler if t["durum"] == "Açık"), key=lambda t: (ONCELIK[t["acik_durum"]], -t["tutar"], t["tarih"]))
    temsilciler = {}
    for t in teklifler:
        temsilciler.setdefault(katla(t["temsilci"]), []).append(t)
    kayip = {}
    for t in teklifler:
        if t["durum"] == "Kaybedildi":
            n = kayip.setdefault(katla(t["neden"]), {"neden": t["neden"] or "(belirtilmemiş)", "adet": 0, "tutar": Z})
            n["adet"] += 1
            n["tutar"] += t["tutar"]
    toplam_kayip = sum(n["adet"] for n in kayip.values())
    for n in kayip.values():
        n["pay"] = _oran(n["adet"], toplam_kayip)
    return {"teklifler": teklifler, "acik": acik, "ozet": _ozet(teklifler), "temsilciler": [{"temsilci": g[0]["temsilci"], **_ozet(g)} for _, g in sorted(temsilciler.items())],
            "dilimler": [{"dilim": e, "adet": sum(1 for t in acik if t["dilim"] == e), "tutar": sum((t["tutar"] for t in acik if t["dilim"] == e), Z)} for e in etiketler],
            "kayip": sorted(kayip.values(), key=lambda n: (-n["adet"], -n["tutar"], katla(n["neden"]))), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, bugun: date) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Açık Teklifler"
    _baslik(ws, ["Durum", "Teklif No", "Teklif Tarihi", "Müşteri", "Tutar", "Bekleme (gün)", "Geçerlilik", "Son Takip", "Takipsiz (gün)", "Temsilci", "Not (siz doldurun)"], (22, 13, 12, 28, 15, 10, 12, 12, 10, 20, 34))
    for t in r["acik"]:
        _satir(ws, [t["acik_durum"], t["no"], t["tarih"], t["musteri"], t["tutar"], t["bekleme"], t["gecerlilik"], t["takip"], t["takipsiz"], t["temsilci"], None], para_s=(5,), tarih_s=(3, 7, 8), dolgu={1: ACIK_RENK[t["acik_durum"]]})
    ws = wb.create_sheet("Özet")
    o = r["ozet"]
    _baslik(ws, ["Gösterge", "Değer"], (44, 18))
    ws.append([f"Durum tarihi: {bugun:%d.%m.%Y}"])
    for etiket, deger, para_mi in (("Teklif sayısı", o["adet"], False), ("Açık teklif", o["acik"], False), ("Açık teklif tutarı", o["acik_tutar"], True), ("Kazanılan", o["kazanilan"], False), ("Kazanılan tutar", o["kazanilan_tutar"], True),
                                   ("Kaybedilen", o["kaybedilen"], False), ("Kaybedilen tutar", o["kaybedilen_tutar"], True), ("İptal", o["iptal"], False), ("Dönüşüm oranı % (adet)", o["donusum"], True),
                                   ("Dönüşüm oranı % (tutar)", o["donusum_tutar"], True), ("Ortalama sonuçlanma süresi (gün)", o["ort_sure"], True)):
        _satir(ws, [etiket, deger], para_s=(2,) if para_mi else ())
    ws.append([])
    _baslik(ws, ["Açık tekliflerin yaşı (gün)", "Adet", "Tutar"], dondur="A2")
    for d in r["dilimler"]:
        _satir(ws, [d["dilim"], d["adet"], d["tutar"]], para_s=(3,))
    ws.column_dimensions["C"].width = 18
    ws = wb.create_sheet("Temsilciler")
    _baslik(ws, ["Temsilci", "Teklif", "Açık", "Açık Tutar", "Kazanılan", "Kazanılan Tutar", "Kaybedilen", "Kaybedilen Tutar", "İptal", "Dönüşüm % (adet)", "Dönüşüm % (tutar)", "Ort. Sonuçlanma (gün)"],
            (24, 8, 8, 15, 10, 15, 10, 15, 8, 11, 11, 12))
    for s in r["temsilciler"]:
        _satir(ws, [s["temsilci"], s["adet"], s["acik"], s["acik_tutar"], s["kazanilan"], s["kazanilan_tutar"], s["kaybedilen"], s["kaybedilen_tutar"], s["iptal"], s["donusum"], s["donusum_tutar"], s["ort_sure"]], para_s=(4, 6, 8, 10, 11, 12))
    ws = wb.create_sheet("Kayıp Nedenleri")
    _baslik(ws, ["Neden", "Adet", "Tutar", "Pay % (adet)"], (32, 8, 16, 11))
    for n in r["kayip"]:
        _satir(ws, [n["neden"], n["adet"], n["tutar"], n["pay"]], para_s=(3, 4))
    ws = wb.create_sheet("Tüm Teklifler")
    _baslik(ws, ["Teklif No", "Teklif Tarihi", "Müşteri", "Tutar", "Durum", "Geçerlilik", "Son Takip", "Sonuç Tarihi", "Sonuçlanma (gün)", "Temsilci", "Kayıp Nedeni"], (13, 12, 28, 15, 12, 12, 12, 12, 11, 20, 24))
    for t in r["teklifler"]:
        _satir(ws, [t["no"], t["tarih"], t["musteri"], t["tutar"], t["durum"], t["gecerlilik"], t["takip"], t["sonuc"], t["sure"], t["temsilci"], t["neden"]], para_s=(4,), tarih_s=(2, 6, 7, 8),
               dolgu={5: {"Kazanıldı": YESIL, "Kaybedildi": KIRMIZI}.get(t["durum"])})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, takip_gun: int = 7, gecerlilik_gun: int = 5, dilimler: tuple[int, ...] = (7, 15, 30)) -> dict:
    teklifler, uy = oku(yol)
    r = degerlendir(teklifler, bugun, takip_gun, gecerlilik_gun, dilimler)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, bugun)
    return r


def dilim_coz(s: str) -> tuple[int, ...] | None:
    try:
        d = tuple(int(x) for x in re.split(r"[;,\s]+", s.strip()) if x)
    except ValueError:
        return None
    return d if d and d[0] >= 0 and all(a < b for a, b in zip(d, d[1:])) else None


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Tekliflerin durumunu, bekleme süresini ve dönüşüm oranını izler.")
    p.add_argument("--teklifler", type=Path, help="Teklif No, Teklif Tarihi, Müşteri, Tutar, Durum, Geçerlilik Tarihi, Son Takip Tarihi, Sonuç Tarihi, Temsilci, Kayıp Nedeni")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--takip", type=int, default=7, help="Bu kadar gündür aranmayan açık teklif için takip uyarısı (varsayılan 7)")
    p.add_argument("--gecerlilik", type=int, default=5, help="Geçerliliğin dolmasına bu kadar gün kala uyarılır (varsayılan 5)")
    p.add_argument("--dilimler", default="7;15;30", help="Açık tekliflerin yaş dilimi sınırları, gün (varsayılan 7;15;30)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "teklif_takip_listesi.xlsx")
    a = p.parse_args(argv)
    ornek = a.teklifler is None
    yol = ORNEK / "teklifler.csv" if ornek else a.teklifler
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    dilimler = dilim_coz(a.dilimler)
    if bugun is None or a.takip < 1 or a.gecerlilik < 0 or dilimler is None:
        print("[X] --tarih GG.AA.YYYY; --takip en az 1; --gecerlilik sıfır veya pozitif; --dilimler artan tam sayılar (örn. 7;15;30) olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, bugun, a.takip, a.gecerlilik, dilimler)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {o['adet']} teklif · açık {o['acik']} ({o['acik_tutar']:,.2f}) · kazanılan {o['kazanilan']}, kaybedilen {o['kaybedilen']}" + (f" · dönüşüm %{o['donusum']} (tutar %{o['donusum_tutar']})" if o["donusum"] is not None else ""))
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
