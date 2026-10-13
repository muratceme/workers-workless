import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)


def t(no, once, durum="Açık", tutar=100, gecerlilik=None, takip=None, sonuc=None, temsilci="A", neden=""):
    """once: teklif kaç gün önce verildi; gecerlilik: bugünden kaç gün sonra; takip: kaç gün önce; sonuc: tekliften kaç gün sonra."""
    g = lambda n: None if n is None else BUGUN + timedelta(days=n)   # noqa: E731
    tarih = BUGUN - timedelta(days=once)
    return {"no": no, "tarih": tarih, "musteri": "M", "tutar": D(tutar), "durum": durum, "gecerlilik": g(gecerlilik), "takip": g(-takip) if takip is not None else None,
            "sonuc": tarih + timedelta(days=sonuc) if sonuc is not None else None, "temsilci": temsilci, "neden": neden}


class DegerlendirTesti(unittest.TestCase):
    def test_acik_teklifler(self):
        l = [t("dün", 3, gecerlilik=-1), t("bugün", 3, gecerlilik=0), t("beş", 3, gecerlilik=5), t("altı", 3, gecerlilik=6), t("takip6", 20, takip=6), t("takip7", 20, takip=7, tutar=50), t("hiç7", 7, tutar=900),
             t("hiç6", 6), t("ikisi", 30, gecerlilik=-2, takip=25), t("ileri", -2), t("takip_once", 3, takip=10)]
        r = main.degerlendir(l, BUGUN)
        self.assertEqual([(x["no"], x["acik_durum"], x["bekleme"], x["takipsiz"], x["dilim"]) for x in r["teklifler"]],
                         [("dün", "Geçerliliği doldu", 3, 3, "0–7"), ("bugün", "Geçerlilik yaklaşıyor", 3, 3, "0–7"), ("beş", "Geçerlilik yaklaşıyor", 3, 3, "0–7"), ("altı", "Bekliyor", 3, 3, "0–7"),
                          ("takip6", "Bekliyor", 20, 6, "16–30"), ("takip7", "Takip gerekli", 20, 7, "16–30"), ("hiç7", "Takip gerekli", 7, 7, "0–7"), ("hiç6", "Bekliyor", 6, 6, "0–7"),
                          ("ikisi", "Geçerliliği doldu", 30, 25, "16–30"), ("ileri", "Bekliyor", 0, 0, "0–7"), ("takip_once", "Bekliyor", 3, 3, "0–7")])
        self.assertEqual([x["no"] for x in r["acik"]], ["ikisi", "dün", "bugün", "beş", "hiç7", "takip7", "takip6", "hiç6", "altı", "takip_once", "ileri"])
        self.assertEqual([(d["dilim"], d["adet"], d["tutar"]) for d in r["dilimler"]], [("0–7", 8, D(1600)), ("8–15", 0, D(0)), ("16–30", 3, D(250)), ("31+", 0, D(0))])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         sorted([("Yüksek", "Geçerliliği dolmuş açık teklif", "dün · M"), ("Yüksek", "Geçerliliği dolmuş açık teklif", "ikisi · M"), ("Orta", "Geçerlilik süresi doluyor", "bugün · M"), ("Orta", "Geçerlilik süresi doluyor", "beş · M"),
                                 ("Orta", "Takip gerekli", "takip7 · M"), ("Orta", "Takip gerekli", "hiç7 · M"), ("Orta", "Teklif tarihi ileride", "ileri · M")]))
        r = main.degerlendir([t("a", 8, gecerlilik=3, takip=4), t("b", 16), t("c", 31), t("d", 61)], BUGUN, takip_gun=4, gecerlilik_gun=2, dilimler=(15, 60))
        self.assertEqual(([(x["acik_durum"], x["dilim"]) for x in r["teklifler"]], [d["dilim"] for d in r["dilimler"]]),
                         ([("Takip gerekli", "0–15"), ("Takip gerekli", "16–60"), ("Takip gerekli", "16–60"), ("Takip gerekli", "61+")], ["0–15", "16–60", "61+"]))

    def test_donusum_ve_ozetler(self):
        l = [t("k1", 40, "Kazanıldı", 300, sonuc=10, temsilci="Ali"), t("k2", 40, "Kazanıldı", 100, sonuc=None, temsilci="ali"), t("y1", 40, "Kaybedildi", 600, sonuc=21, temsilci="Ali", neden="Fiyat"),
             t("y2", 40, "Kaybedildi", 200, sonuc=-1, temsilci="Ece", neden="fiyat"), t("y3", 40, "Kaybedildi", 50, sonuc=0, temsilci="Ece"), t("i", 40, "İptal", 9999, temsilci="Ece"), t("a", 1, tutar=70, temsilci="Can")]
        r = main.degerlendir(l, BUGUN)
        self.assertEqual(r["ozet"], {"adet": 7, "acik": 1, "acik_tutar": D(70), "kazanilan": 2, "kazanilan_tutar": D(400), "kaybedilen": 3, "kaybedilen_tutar": D(850), "iptal": 1, "donusum": D("40.00"), "donusum_tutar": D("32.00"),
                                     "ort_sure": D("10.33")})
        self.assertEqual([(s["temsilci"], s["adet"], s["acik"], s["kazanilan"], s["kaybedilen"], s["iptal"], s["donusum"], s["donusum_tutar"], s["ort_sure"]) for s in r["temsilciler"]],
                         [("Ali", 3, 0, 2, 1, 0, D("66.67"), D("40.00"), D("15.50")), ("Can", 1, 1, 0, 0, 0, None, None, None), ("Ece", 3, 0, 0, 2, 1, D("0.00"), D("0.00"), D("0.00"))])
        self.assertEqual(r["kayip"], [{"neden": "Fiyat", "adet": 2, "tutar": D(800), "pay": D("66.67")}, {"neden": "(belirtilmemiş)", "adet": 1, "tutar": D(50), "pay": D("33.33")}])
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]), [("Kayıp nedeni yazılmamış", "y3 · M"), ("Sonuç tarihi yok", "k2 · M"), ("Sonuç tarihi yok", "y2 · M")])
        self.assertEqual([x["sure"] for x in r["teklifler"]], [10, None, 21, None, 0, None, None])
        with self.assertRaises(ValueError):
            main.degerlendir([], BUGUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "teklifler.csv", cikti, BUGUN)
            self.assertEqual(r["ozet"], {"adet": 13, "acik": 6, "acik_tutar": D(2595000), "kazanilan": 3, "kazanilan_tutar": D(1915000), "kaybedilen": 3, "kaybedilen_tutar": D(805000), "iptal": 1, "donusum": D("50.00"),
                                         "donusum_tutar": D("70.40"), "ort_sure": D("22.60")})
            self.assertEqual([(x["no"], x["acik_durum"], x["bekleme"], x["takipsiz"]) for x in r["acik"]],
                             [("T-103", "Geçerliliği doldu", 32, 22), ("T-104", "Geçerlilik yaklaşıyor", 27, 4), ("T-105", "Takip gerekli", 22, 17), ("T-107", "Takip gerekli", 7, 7), ("T-106", "Bekliyor", 14, 3), ("T-112", "Bekliyor", 4, 4)])
            self.assertEqual([(s["temsilci"], s["adet"], s["acik"], s["acik_tutar"], s["kazanilan_tutar"], s["kaybedilen_tutar"], s["iptal"], s["donusum"], s["donusum_tutar"], s["ort_sure"]) for s in r["temsilciler"]],
                             [("Burak Şen", 5, 2, D(1330000), D(480000), D(150000), 1, D("50.00"), D("76.19"), D("19.50")), ("Deniz Acar", 5, 2, D(530000), D(185000), D(655000), 0, D("33.33"), D("22.02"), D("25.00")),
                              ("Ece Yurt", 3, 2, D(735000), D(1250000), D(0), 0, D("100.00"), D("100.00"), D("24.00"))])
            self.assertEqual([(d["dilim"], d["adet"], d["tutar"]) for d in r["dilimler"]], [("0–7", 2, D(525000)), ("8–15", 1, D(640000)), ("16–30", 2, D(1110000)), ("31+", 1, D(320000))])
            self.assertEqual(r["kayip"], [{"neden": "Fiyat", "adet": 2, "tutar": D(425000), "pay": D("66.67")}, {"neden": "Teslim Süresi", "adet": 1, "tutar": D(380000), "pay": D("33.33")}])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Yüksek", "Geçerliliği dolmuş açık teklif", "T-103 · Gama Tekstil"), ("Orta", "Geçerlilik süresi doluyor", "T-104 · Delta Yapı"), ("Orta", "Takip gerekli", "T-105 · Epsilon Lojistik"),
                              ("Orta", "Takip gerekli", "T-107 · Eta Plastik"), ("Bilgi", "Sonuç tarihi yok", "T-113 · Nü Ambalaj")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Açık Teklifler", "Özet", "Temsilciler", "Kayıp Nedenleri", "Tüm Teklifler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "t.csv"
            y.write_text("No;Tarih;Firma;Toplam;Durum;Opsiyon Tarihi;Son Görüşme;Kapanış Tarihi;Sorumlu;Ret Nedeni\nT1;01.10.2026;A;1.250,50;SİPARİŞE DÖNÜŞTÜ;;;05.10.2026;;\nt1;02.10.2026;A;5;;10.10.2026;03.10.2026;;Ali;\n"
                         "T2;x;A;5;;;;;;\nT3;02.10.2026;A;5;Belirsiz;;;;;\nT4;02.10.2026;A;-5;;;;;;\n;02.10.2026;A;5;;;;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["tarih"], x["tutar"], x["durum"], x["gecerlilik"], x["takip"], x["sonuc"], x["temsilci"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("T1", date(2026, 10, 1), D("1250.50"), "Kazanıldı", None, None, date(2026, 10, 5), "(belirtilmemiş)"), ("t1", date(2026, 10, 2), D(5), "Açık", date(2026, 10, 10), date(2026, 10, 3), None, "Ali")],
                              [("Mükerrer teklif numarası", "t1"), ("Teklif okunamadı", "T2"), ("Teklif okunamadı", "T3"), ("Teklif okunamadı", "T4")]))
            self.assertEqual([main.dilim_coz(s) for s in ("7;15;30", "10, 20", "30", "15;7", "7;7", "x", "", "-1;5")], [(7, 15, 30), (10, 20), (30,), None, None, None, None, None])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--teklifler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "12.10.2026", "--takip", "5", "--gecerlilik", "3", "--dilimler", "10;20;45"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + x) for x in (["--tarih", "x"], ["--takip", "0"], ["--gecerlilik", "-1"], ["--dilimler", "30;15"])], [2] * 4)
            self.assertEqual(main.main(["--teklifler", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Teklif No;Tutar\nA;1\n", "Teklif No;Tarih;Tutar\n;01.10.2026;1\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
