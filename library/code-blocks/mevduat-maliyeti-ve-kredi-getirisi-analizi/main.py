"""
Mevduat Maliyeti ve Kredi Getirisi Analizi — Gerisi Bende kod bloğu
Hazine ve Aktif-Pasif Yönetimi › Aktif-Pasif Yönetimi Uzmanı

Ürün bazında ağırlıklı ortalama mevduat maliyetini ve kredi getirisini hesaplar; marj analizini yapar.
  Yıllık faiz tutarı     = bakiye × faiz oranı ÷ 100
  Ağırlıklı ortalama %   = Σ (bakiye × faiz oranı) ÷ Σ bakiye   (tür ve para birimi bazında)
  Vadesiz hariç maliyet  = faiz oranı sıfırdan büyük mevduat ürünleri üzerinden aynı hesap
  Ürünün katkısı (puan)  = ürünün gruptaki payı × faiz oranı   (katkıların toplamı = ağırlıklı ortalama)
  Marj (puan)            = ağırlıklı ortalama kredi getirisi − ağırlıklı ortalama mevduat maliyeti   (aynı para biriminde)
  Kredi / mevduat %      = kredi bakiyesi ÷ mevduat bakiyesi × 100
  Net faiz (basit)       = kredilerin yıllık faiz tutarı − mevduatın yıllık faiz tutarı
Oranlar yazdığınız basit yıllık oranlardır. Zorunlu karşılık, vergi ve fon kesintileri, bileşik faiz ve vade farkı hesaba katılmaz; para birimleri birbirine çevrilmez.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
SUTUN = {"ad": ("urun", "urun adi", "kalem"), "tur": ("tur", "urun turu", "sinif"), "pb": ("para birimi", "doviz", "doviz cinsi", "pb"), "bakiye": ("bakiye", "tutar", "ortalama bakiye"), "oran": ("faiz orani", "faiz", "oran", "getiri")}
TUR = {"mevduat": "Mevduat", "katilim fonu": "Mevduat", "kaynak": "Mevduat", "kredi": "Kredi", "krediler": "Kredi", "plasman": "Kredi"}
Z, YUZ = Decimal(0), Decimal(100)


def olcu(x) -> Decimal | None:
    """Oran okur: virgül ya da tek nokta ondalık ayırıcıdır (3,5 = 3.5)."""
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x or ""))
    try:
        return Decimal(s.replace(".", "").replace(",", ".") if "," in s else s) if s else None
    except InvalidOperation:
        return None


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("ad", "tur", "bakiye", "oran")):
        ad = metin(r["ad"])
        if not ad:
            continue
        tur, bakiye, oran = TUR.get(katla(r["tur"])), para(r["bakiye"]), olcu(r["oran"])
        if tur is None or bakiye is None or bakiye < 0 or oran is None or oran < 0:
            uy.append(uyari("Yüksek", "Ürün okunamadı", ad, f"Satır {r['_satir']}: tür '{metin(r['tur'])}' (Mevduat / Kredi olmalı), bakiye '{metin(r['bakiye'])}', faiz oranı '{metin(r['oran'])}'. Vadesiz için oranı 0 yazın. Satır alınmadı."))
            continue
        pb = katla(r["pb"]).upper() or "TRY"
        l.append({"ad": ad, "tur": tur, "pb": "TRY" if pb == "TL" else pb, "bakiye": bakiye, "oran": oran})
    return l, uy


def _ort(liste: list[dict]) -> Decimal | None:
    b = sum((u["bakiye"] for u in liste), Z)
    return sum((u["bakiye"] * u["oran"] for u in liste), Z) / b if b else None


def _yv(x: Decimal | None) -> Decimal | None:
    return yuvarla(x) if x is not None else None


def analiz(urunler: list[dict]) -> dict:
    if not urunler:
        raise ValueError("Analiz edilecek ürün yok.")
    uy, gruplar, marjlar = [], [], []
    for pb in sorted({u["pb"] for u in urunler}, key=lambda p: (p != "TRY", p)):
        ham = {}
        for tur in ("Mevduat", "Kredi"):
            liste = [u for u in urunler if u["pb"] == pb and u["tur"] == tur]
            bakiye, ort = sum((u["bakiye"] for u in liste), Z), _ort(liste)
            ham[tur] = (bakiye, ort, sum((u["bakiye"] * u["oran"] / YUZ for u in liste), Z))
            for u in liste:
                u["faiz"] = yuvarla(u["bakiye"] * u["oran"] / YUZ)
                u["pay"] = yuvarla(u["bakiye"] * YUZ / bakiye) if bakiye else None
                u["katki"] = yuvarla(u["bakiye"] * u["oran"] / bakiye) if bakiye else None
            if liste:
                gruplar.append({"tur": tur, "pb": pb, "urun": len(liste), "bakiye": bakiye, "faiz": yuvarla(ham[tur][2]), "oran": _yv(ort), "vadesiz_haric": _yv(_ort([u for u in liste if u["oran"] > 0])) if tur == "Mevduat" else None})
        (mb, mo, mf), (kb, ko, kf) = ham["Mevduat"], ham["Kredi"]
        m = {"pb": pb, "mevduat": mb, "kredi": kb, "maliyet": _yv(mo), "getiri": _yv(ko), "marj": _yv(ko - mo) if mo is not None and ko is not None else None, "kredi_mevduat": yuvarla(kb * YUZ / mb) if mb else None, "net_faiz": yuvarla(kf - mf)}
        marjlar.append(m)
        if m["marj"] is not None and m["marj"] < 0:
            uy.append(uyari("Yüksek", "Marj eksi", pb, f"Ağırlıklı ortalama kredi getirisi %{m['getiri']}, mevduat maliyeti %{m['maliyet']}; marj {m['marj']} puan."))
        if mo is not None and ko is not None:
            for u in urunler:
                if u["pb"] == pb and u["tur"] == "Kredi" and u["oran"] < mo:
                    uy.append(uyari("Orta", "Kredi getirisi mevduat maliyetinin altında", f"{u['ad']} ({pb})", f"Ürünün faiz oranı %{u['oran']:f}, ağırlıklı ortalama mevduat maliyeti %{m['maliyet']}. Bakiye {u['bakiye']:,.2f}."))
                if u["pb"] == pb and u["tur"] == "Mevduat" and u["oran"] > ko:
                    uy.append(uyari("Orta", "Mevduat maliyeti kredi getirisinin üstünde", f"{u['ad']} ({pb})", f"Ürünün faiz oranı %{u['oran']:f}, ağırlıklı ortalama kredi getirisi %{m['getiri']}. Bakiye {u['bakiye']:,.2f}."))
        if not mb or not kb:
            uy.append(uyari("Bilgi", "Marj hesaplanamadı", pb, f"Bu para biriminde {'mevduat' if not mb else 'kredi'} bakiyesi yok; marj hesaplanmadı."))
    sira = {"Mevduat": 0, "Kredi": 1}
    urunler.sort(key=lambda u: (u["pb"] != "TRY", u["pb"], sira[u["tur"]], -u["bakiye"], katla(u["ad"])))
    return {"urunler": urunler, "gruplar": gruplar, "marjlar": marjlar, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Marj"
    _baslik(ws, ["Para Birimi", "Mevduat Bakiyesi", "Kredi Bakiyesi", "Kredi / Mevduat %", "Mevduat Maliyeti %", "Kredi Getirisi %", "Marj (puan)", "Net Faiz (yıllık, basit)"], (10, 19, 19, 11, 12, 12, 11, 19))
    for m in r["marjlar"]:
        _satir(ws, [m["pb"], m["mevduat"], m["kredi"], m["kredi_mevduat"], m["maliyet"], m["getiri"], m["marj"], m["net_faiz"]], para_s=(2, 3, 4, 5, 6, 7, 8), dolgu={7: None if m["marj"] is None else KIRMIZI if m["marj"] < 0 else YESIL})
    ws.append([])
    _baslik(ws, ["Tür", "Para Birimi", "Ürün Sayısı", "Bakiye", "Yıllık Faiz Tutarı", "Ağırlıklı Ortalama %", "Vadesiz Hariç %"], dondur="A2")
    for g in r["gruplar"]:
        _satir(ws, [g["tur"], g["pb"], g["urun"], g["bakiye"], g["faiz"], g["oran"], g["vadesiz_haric"]], para_s=(4, 5, 6, 7))
    ws.append([])
    ws.append(["Marj = kredi getirisi − mevduat maliyeti (aynı para biriminde, basit yıllık oranlarla). Zorunlu karşılık, kesintiler, operasyonel maliyet ve kredi zararı dahil değildir."])
    ws = wb.create_sheet("Ürünler")
    _baslik(ws, ["Para Birimi", "Tür", "Ürün", "Bakiye", "Gruptaki Payı %", "Faiz Oranı %", "Ortalamaya Katkı (puan)", "Yıllık Faiz Tutarı"], (10, 10, 34, 19, 11, 10, 13, 19))
    for u in r["urunler"]:
        _satir(ws, [u["pb"], u["tur"], u["ad"], u["bakiye"], u["pay"], u["oran"], u["katki"], u["faiz"]], para_s=(4, 5, 6, 7, 8))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path) -> dict:
    urunler, uy = oku(yol)
    r = analiz(urunler)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Ürün bazında ağırlıklı ortalama mevduat maliyetini, kredi getirisini ve marjı hesaplar.")
    p.add_argument("--urunler", type=Path, help="Ürün, Tür (Mevduat / Kredi), Para Birimi, Bakiye, Faiz Oranı")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "mevduat_maliyeti_kredi_getirisi.xlsx")
    a = p.parse_args(argv)
    yol = a.urunler or ORNEK / "urunler.csv"
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    for m in r["marjlar"]:
        print(f"[OK] {m['pb']}: mevduat maliyeti " + (f"%{m['maliyet']}" if m["maliyet"] is not None else "—") + " · kredi getirisi " + (f"%{m['getiri']}" if m["getiri"] is not None else "—")
              + (f" · marj {m['marj']} puan" if m["marj"] is not None else "") + (f" · kredi/mevduat %{m['kredi_mevduat']}" if m["kredi_mevduat"] is not None else ""))
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
