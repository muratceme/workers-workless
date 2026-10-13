import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
M, K = "Mevduat", "Kredi"


def u(ad, tur, bakiye, oran, pb="TRY"):
    return {"ad": ad, "tur": tur, "pb": pb, "bakiye": D(str(bakiye)), "oran": D(str(oran))}


class AnalizTesti(unittest.TestCase):
    def test_ortalama_ve_marj(self):
        r = main.analiz([u("vadesiz", M, 400, 0), u("v1", M, 600, 42), u("v2", M, 800, 45), u("v3", M, 200, 40), u("k1", K, 900, 52), u("k2", K, 600, "50.5")])
        self.assertEqual(r["gruplar"], [{"tur": M, "pb": "TRY", "urun": 4, "bakiye": D(2000), "faiz": D("692.00"), "oran": D("34.60"), "vadesiz_haric": D("43.25")},
                                        {"tur": K, "pb": "TRY", "urun": 2, "bakiye": D(1500), "faiz": D("771.00"), "oran": D("51.40"), "vadesiz_haric": None}])
        self.assertEqual(r["marjlar"], [{"pb": "TRY", "mevduat": D(2000), "kredi": D(1500), "maliyet": D("34.60"), "getiri": D("51.40"), "marj": D("16.80"), "kredi_mevduat": D("75.00"), "net_faiz": D("79.00")}])
        self.assertEqual([(x["ad"], x["pay"], x["katki"], x["faiz"]) for x in r["urunler"]],
                         [("v2", D("40.00"), D("18.00"), D("360.00")), ("v1", D("30.00"), D("12.60"), D("252.00")), ("vadesiz", D("20.00"), D("0.00"), D("0.00")), ("v3", D("10.00"), D("4.00"), D("80.00")),
                          ("k1", D("60.00"), D("31.20"), D("468.00")), ("k2", D("40.00"), D("20.20"), D("303.00"))])
        self.assertEqual((sum(x["katki"] for x in r["urunler"][:4]), r["uyarilar"]), (D("34.60"), []))
        r = main.analiz([u("m", M, 3, 1), u("m2", M, 3, 2), u("k", K, 1, "1.5")])
        self.assertEqual((r["gruplar"][0]["oran"], r["marjlar"][0]["marj"], r["marjlar"][0]["kredi_mevduat"]), (D("1.50"), D("0.00"), D("16.67")))
        with self.assertRaises(ValueError):
            main.analiz([])

    def test_uyarilar_ve_para_birimi(self):
        l = [u("pahalı mevduat", M, 100, 60), u("mevduat", M, 100, 40), u("ucuz kredi", K, 100, 30), u("kredi", K, 100, 50), u("usd mevduat", M, 50, 2, "USD"), u("eur kredi", K, 70, 9, "EUR"), u("usd vadesiz", M, 50, 0, "USD"),
             u("gbp m", M, 0, 5, "GBP"), u("gbp k", K, 10, 4, "GBP")]
        r = main.analiz(l)
        self.assertEqual([(m["pb"], m["maliyet"], m["getiri"], m["marj"], m["kredi_mevduat"], m["net_faiz"]) for m in r["marjlar"]],
                         [("TRY", D("50.00"), D("40.00"), D("-10.00"), D("100.00"), D("-20.00")), ("EUR", None, D("9.00"), None, None, D("6.30")), ("GBP", None, D("4.00"), None, None, D("0.40")), ("USD", D("1.00"), None, None, D("0.00"), D("-1.00"))])
        self.assertEqual(sorted((w["onem"], w["tur"], w["kim"]) for w in r["uyarilar"]),
                         sorted([("Yüksek", "Marj eksi", "TRY"), ("Orta", "Kredi getirisi mevduat maliyetinin altında", "ucuz kredi (TRY)"), ("Orta", "Mevduat maliyeti kredi getirisinin üstünde", "pahalı mevduat (TRY)"),
                                 ("Bilgi", "Marj hesaplanamadı", "EUR"), ("Bilgi", "Marj hesaplanamadı", "GBP"), ("Bilgi", "Marj hesaplanamadı", "USD")]))
        self.assertEqual(([(g["tur"], g["pb"], g["oran"], g["vadesiz_haric"]) for g in r["gruplar"]][-1], [x["pb"] for x in r["urunler"]][:5], [x["pay"] for x in r["urunler"] if x["ad"] == "gbp m"]),
                         ((M, "USD", D("1.00"), D("2.00")), ["TRY"] * 4 + ["EUR"], [None]))


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "urunler.csv", cikti)
            m = 1000000
            self.assertEqual(r["gruplar"], [{"tur": M, "pb": "TRY", "urun": 4, "bakiye": D(2000 * m), "faiz": D(692 * m), "oran": D("34.60"), "vadesiz_haric": D("43.25")},
                                            {"tur": K, "pb": "TRY", "urun": 4, "bakiye": D(1500 * m), "faiz": D(769 * m), "oran": D("51.27"), "vadesiz_haric": None},
                                            {"tur": M, "pb": "USD", "urun": 2, "bakiye": D(40 * m), "faiz": D(1050000), "oran": D("2.63"), "vadesiz_haric": D("3.50")},
                                            {"tur": K, "pb": "USD", "urun": 1, "bakiye": D(25 * m), "faiz": D(2125000), "oran": D("8.50"), "vadesiz_haric": None}])
            self.assertEqual(r["marjlar"], [{"pb": "TRY", "mevduat": D(2000 * m), "kredi": D(1500 * m), "maliyet": D("34.60"), "getiri": D("51.27"), "marj": D("16.67"), "kredi_mevduat": D("75.00"), "net_faiz": D(77 * m)},
                                            {"pb": "USD", "mevduat": D(40 * m), "kredi": D(25 * m), "maliyet": D("2.63"), "getiri": D("8.50"), "marj": D("5.88"), "kredi_mevduat": D("62.50"), "net_faiz": D(1075000)}])
            self.assertEqual([(x["ad"], x["pay"], x["katki"]) for x in r["urunler"] if x["tur"] == K and x["pb"] == "TRY"],
                             [("Ticari Krediler", D("60.00"), D("31.20")), ("Bireysel İhtiyaç Kredisi", D("20.00"), D("12.40")), ("Konut Kredisi", D("13.33"), D("4.00")), ("Kredi Kartı", D("6.67"), D("3.67"))])
            self.assertEqual([(w["onem"], w["tur"], w["kim"]) for w in r["uyarilar"]], [("Orta", "Kredi getirisi mevduat maliyetinin altında", "Konut Kredisi (TRY)")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Marj", "Ürünler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "u.csv"
            y.write_text("Kalem;Sınıf;Döviz;Ortalama Bakiye;Getiri\nVadeli;KAYNAK;tl;1.250.000,50;%42,5\nTicari;Plasman;usd;5;8.5\nX;Fon;;5;5\nY;Kredi;;x;5\nZ;Kredi;;5;\nW;Kredi;;5;-1\n;Kredi;;5;5\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual((l, [(w["tur"], w["kim"]) for w in uy]),
                             ([{"ad": "Vadeli", "tur": M, "pb": "TRY", "bakiye": D("1250000.50"), "oran": D("42.5")}, {"ad": "Ticari", "tur": K, "pb": "USD", "bakiye": D(5), "oran": D("8.5")}], [("Ürün okunamadı", n) for n in "XYZW"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--urunler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(["--urunler", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Ürün;Tür;Bakiye\nA;Kredi;1\n", "Ürün;Tür;Bakiye;Faiz Oranı\n;Kredi;1;1\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
