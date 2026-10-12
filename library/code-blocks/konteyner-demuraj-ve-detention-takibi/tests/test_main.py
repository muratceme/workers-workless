import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
RAPOR = date(2026, 10, 12)


def tf(tur, bas, bit, ucret, hat="", tip=""):
    return {"hat": hat, "tip": tip, "tur": tur, "bas": bas, "bit": bit, "ucret": D(ucret), "para": "USD"}


TARIFE = [tf("demuraj", 1, 7, 0), tf("demuraj", 8, 14, 40), tf("demuraj", 15, None, 80), tf("detention", 1, 10, 0), tf("detention", 11, 20, 30), tf("detention", 21, None, 60)]


def kn(tahliye, cikis=None, iade=None, no="K", hat="", tip=""):
    t = lambda g: None if g is None else RAPOR + timedelta(days=g)   # noqa: E731
    return {"no": no, "tip": tip, "hat": hat, "ref": "", "tahliye": t(tahliye), "cikis": t(cikis), "iade": t(iade)}


def tek(*a, tarife=TARIFE, **kw):
    dahil = kw.pop("dahil", True)
    return main.takip([kn(*a, **kw)], tarife, RAPOR, dahil)["konteynerler"][0]


class TakipTesti(unittest.TestCase):
    def test_dilim_ucreti(self):
        dl = TARIFE[:3]
        self.assertEqual([main.ucret(g, dl) for g in (0, 7, 8, 14, 15, 18)], [(D(0), 7, D(0)), (D(0), 7, D(0)), (D(40), 7, D(40)), (D(280), 7, D(40)), (D(360), 7, D(80)), (D(600), 7, D(80))])
        self.assertEqual(main.ucret(3, [tf("demuraj", 1, None, 25)]), (D(75), 0, D(25)))                       # serbest süresi olmayan tarife

    def test_kapali_ve_acik(self):
        k = tek(-11, -7, -3)
        self.assertEqual((k["kalemler"]["demuraj"]["gun"], k["kalemler"]["detention"]["gun"], k["toplam"], k["durum"], k["acik"]), (5, 5, D(0), "Kapalı", ""))
        k = tek(-41, -39, -12)
        self.assertEqual((k["kalemler"]["detention"]["gun"], k["kalemler"]["detention"]["ucret"], k["toplam"]), (28, D(780), D(780)))
        k = tek(-17)                                                                                            # limanda, 18. gün
        d = k["kalemler"]["demuraj"]
        self.assertEqual((d["gun"], d["ucret"], d["kalan"], d["gunluk"], k["durum"], k["acik"], "detention" in k["kalemler"]), (18, D(600), 0, D(80), "Ücret işliyor", "Demuraj", False))
        k = tek(-4)
        d = k["kalemler"]["demuraj"]
        self.assertEqual((d["gun"], d["kalan"], d["son_serbest"], k["durum"]), (5, 2, date(2026, 10, 14), "Serbest süre içinde"))
        k = tek(-22, -11)                                                                                       # limandan çıkmış, iade edilmemiş
        self.assertEqual((k["kalemler"]["demuraj"]["ucret"], k["kalemler"]["detention"]["gun"], k["kalemler"]["detention"]["ucret"], k["toplam"], k["acik"], k["acik_kalem"]["gunluk"]),
                         (D(200), 12, D(60), D(260), "Detention", D(30)))
        self.assertEqual([tek(g)["durum"] for g in (-6, -7)], ["Serbest süre içinde", "Ücret işliyor"])        # 7. gün serbest, 8. gün ücretli
        k = tek(-7, dahil=False)
        self.assertEqual((k["kalemler"]["demuraj"]["gun"], k["durum"], k["kalemler"]["demuraj"]["son_serbest"]), (7, "Serbest süre içinde", RAPOR))

    def test_tarife_secimi_ve_toplam(self):
        t = TARIFE + [tf("demuraj", 1, 3, 0, tip="40hc"), tf("demuraj", 4, None, 100, tip="40hc"), tf("demuraj", 1, None, 10, hat="hat x"), tf("demuraj", 1, None, 500, hat="hat x", tip="40hc")]
        self.assertEqual([tek(-4, -1, -1, tarife=t, hat=h, tip=p)["kalemler"]["demuraj"]["ucret"] for h, p in (("", "20DC"), ("", "40HC"), ("Hat X", "20DC"), ("HAT X", "40hc"))], [D(0), D(100), D(40), D(2000)])
        k = tek(-4, tarife=[tf("detention", 1, None, 5)])
        self.assertEqual((k["durum"], k["kalemler"]["demuraj"]["ucret"]), ("Tarife yok", None))
        r = main.takip([kn(-17, no="A"), kn(-4, no="B"), kn(-22, -11, no="C"), kn(-11, -7, -3, no="D")], TARIFE, RAPOR)
        self.assertEqual(([k["no"] for k in r["konteynerler"]], r["toplam"], r["say"]),
                         (["A", "C", "B", "D"], {"demuraj": D(800), "detention": D(60), "genel": D(860), "gunluk": D(110)}, {"Ücret işliyor": 2, "Serbest süre içinde": 1, "Tarife yok": 0, "Kapalı": 1}))
        with self.assertRaises(ValueError):
            main.takip([], TARIFE, RAPOR)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "d.xlsx"
            r = main.calistir(ORNEK / "konteynerler.csv", ORNEK / "tarife.csv", cikti, RAPOR)
            self.assertEqual({k["no"][-1]: (k["durum"], k["toplam"], k["acik"]) for k in r["konteynerler"]},
                             {"1": ("Kapalı", D(0), ""), "2": ("Ücret işliyor", D(410), "Detention"), "3": ("Ücret işliyor", D(600), "Demuraj"), "4": ("Serbest süre içinde", D(0), "Demuraj"),
                              "5": ("Kapalı", D(780), "")})
            self.assertEqual((r["toplam"], r["uyarilar"]), ({"demuraj": D(950), "detention": D(840), "genel": D(1790), "gunluk": D(110)}, []))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Konteyner Takibi", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            k, t = Path(tmp) / "k.csv", Path(tmp) / "t.csv"
            k.write_text("Konteyner No;Tip;Tahliye Tarihi;Limandan Çıkış;Boş İade\nA;20DC;01.10.2026;yarın;\nB;20DC;dün;;\nC;20DC;05.10.2026;01.10.2026;\nD;20DC;01.10.2026;;05.10.2026\nE;45HC;15.10.2026;;\n",
                         encoding="utf-8")
            t.write_text("Tip;Ücret Türü;Gün Başlangıç;Gün Bitiş;Günlük Ücret;Para Birimi\n20DC;Demuraj;1;5;0;usd\n20DC;DEMURRAGE;6;;50;USD\n20DC;Ardiye;1;;5;USD\n20DC;Detention;5;3;5;USD\n", encoding="utf-8")
            tl, uy = main.tarife_oku(t)
            self.assertEqual(([(x["tip"], x["tur"], x["bas"], x["bit"], x["ucret"]) for x in tl], [u["tur"] for u in uy]), ([("20dc", "demuraj", 1, 5, D(0)), ("20dc", "demuraj", 6, None, D(50))], ["Tarife satırı okunamadı"] * 2))
            kl, uy = main.oku(k)
            self.assertEqual(([(x["no"], x["cikis"]) for x in kl], sorted(u["tur"] for u in uy)), ([("A", None), ("E", None)], ["Tahliye tarihi okunamadı", "Tarih okunamadı", "Tarih sırası hatalı", "İade var, çıkış yok"]))
            c = str(Path(tmp) / "x.xlsx")
            r = main.calistir(k, t, Path(c), RAPOR)
            self.assertEqual(({x["no"]: (x["durum"], x["toplam"]) for x in r["konteynerler"]}, sorted({u["tur"] for u in r["uyarilar"]} - {u["tur"] for u in uy})),
                             ({"A": ("Ücret işliyor", D(350)), "E": ("Tarife yok", D(0))}, ["Tahliye tarihi ileri tarihli", "Tarife bulunamadı", "Tarife satırı okunamadı"]))
            arg = ["--konteynerler", str(k), "--tarife", str(t), "--cikti", c, "--tarih", "12.10.2026"]
            self.assertEqual(main.main(arg + ["--gun-sayimi", "haric"]), 0)
            self.assertEqual(main.main(arg[:-1] + ["dün"]), 2)
            self.assertEqual(main.main(arg[:3] + [str(Path(tmp) / "yok.csv")]), 1)
            t.write_text("Ücret Türü;Gün Başlangıç;Günlük Ücret\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
