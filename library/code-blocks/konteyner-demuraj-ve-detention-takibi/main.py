"""
Konteyner Demuraj ve Detention Takibi — Gerisi Bende kod bloğu
Deniz ve Hava Yolu Operasyon › Denizyolu Operasyon Uzmanı

Konteynerlerin limanda bekleme ve iade sürelerini serbest gün kurallarıyla izleyip oluşacak masrafı hesaplar.
  Demuraj günü   = limandan çıkış tarihi − tahliye tarihi + 1        (konteyner limandayken; çıkış yoksa rapor tarihine kadar)
  Detention günü = boş iade tarihi − limandan çıkış tarihi + 1       (konteyner liman dışındayken; iade yoksa rapor tarihine kadar)
                   --gun-sayimi haric verilirse başlangıç günü sayılmaz (+ 1 yapılmaz)
  Ücret          = tarife dosyasındaki gün dilimlerine göre: dilime düşen gün sayısı × günlük ücret; ücreti 0 olan ilk dilimler serbest süredir
  Tarife seçimi  = konteynerin hattı ve tipine en özel satırlar: hat + tip > yalnız hat > yalnız tip > genel
  Açık konteyner = çıkışı veya iadesi olmayan; kalan serbest gün, son serbest gün ve bugün işleyen günlük ücret gösterilir
Serbest süreler, dilimler, ücretler ve gün sayımı hattın tarifesine ve sözleşmenize göre değişir; kod hiçbir hattın tarifesini bilmez.
İthalat akışı (tahliye → limandan çıkış → boş iade) esas alınmıştır. Demuraj ve detention'ın tek serbest sürede birleştiği tarifeler desteklenmez.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


KONT = {"no": ("konteyner no", "konteyner", "konteyner numarasi"), "tip": ("tip", "konteyner tipi", "tur"), "hat": ("hat", "armator", "tasiyici"), "ref": ("dosya no", "referans", "konsimento no", "musteri"),
        "tahliye": ("tahliye tarihi", "tahliye", "gemiden tahliye", "bosaltma tarihi"), "cikis": ("limandan cikis", "cikis tarihi", "gate out", "limandan cikis tarihi"),
        "iade": ("bos iade", "bos iade tarihi", "iade tarihi", "bos teslim")}
TARIFE = {"hat": KONT["hat"], "tip": KONT["tip"], "tur": ("ucret turu", "masraf turu", "demuraj detention", "tur adi"), "bas": ("gun baslangic", "baslangic gunu", "baslangic", "ilk gun"),
          "bit": ("gun bitis", "bitis gunu", "bitis", "son gun"), "ucret": ("gunluk ucret", "ucret", "gunluk"), "para": ("para birimi", "doviz")}
Z = Decimal(0)


def tarife_oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, TARIFE, ("tur", "bas", "ucret")):
        tur = katla(r["tur"])
        if not tur:
            continue
        tur = "demuraj" if tur.startswith("dem") else "detention" if tur.startswith("det") else None
        bas, bit, ucret = para(r["bas"]), para(r["bit"]), para(r["ucret"])
        if tur is None or bas is None or bas < 1 or ucret is None or ucret < 0 or (bit is not None and bit < bas):
            uy.append(uyari("Yüksek", "Tarife satırı okunamadı", f"{yol.name} satır {r['_satir']}", "Ücret türü Demuraj veya Detention; gün başlangıcı 1 veya büyük; bitiş başlangıçtan küçük olamaz. Satır alınmadı."))
            continue
        l.append({"hat": katla(r["hat"]), "tip": katla(r["tip"]), "tur": tur, "bas": int(bas), "bit": int(bit) if bit is not None else None, "ucret": ucret, "para": metin(r["para"]).upper()})
    if not l:
        raise ValueError(f"{yol.name}: tarife dosyasında okunabilir satır yok.")
    return l, uy


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, KONT, ("no", "tahliye")):
        no = metin(r["no"])
        if not no:
            continue
        kim = f"{no} (satır {r['_satir']})"
        t = {a: tarih(r[a]) for a in ("tahliye", "cikis", "iade")}
        if t["tahliye"] is None:
            uy.append(uyari("Yüksek", "Tahliye tarihi okunamadı", kim, f"'{metin(r['tahliye'])}'. Konteyner alınmadı."))
            continue
        for a, ad in (("cikis", "Limandan çıkış"), ("iade", "Boş iade")):
            if t[a] is None and metin(r[a]):
                uy.append(uyari("Yüksek", "Tarih okunamadı", kim, f"{ad} '{metin(r[a])}'. Boş sayıldı; konteyner açık görünür."))
        if t["cikis"] and t["cikis"] < t["tahliye"] or (t["iade"] and t["cikis"] and t["iade"] < t["cikis"]):
            uy.append(uyari("Yüksek", "Tarih sırası hatalı", kim, "Sıra tahliye → limandan çıkış → boş iade olmalı. Konteyner alınmadı."))
            continue
        if t["iade"] and not t["cikis"]:
            uy.append(uyari("Yüksek", "İade var, çıkış yok", kim, "Boş iade tarihi yazılmış ama limandan çıkış tarihi boş. Konteyner alınmadı."))
            continue
        l.append({"no": no, "tip": metin(r["tip"]), "hat": metin(r["hat"]), "ref": metin(r["ref"])} | t)
    return l, uy


def dilimler(tarife: list[dict], tur: str, hat: str, tip: str) -> list[dict]:
    aday = [t for t in tarife if t["tur"] == tur and t["hat"] in ("", katla(hat)) and t["tip"] in ("", katla(tip))]
    if not aday:
        return []
    en = max((bool(t["hat"]) * 2 + bool(t["tip"])) for t in aday)
    return sorted((t for t in aday if bool(t["hat"]) * 2 + bool(t["tip"]) == en), key=lambda t: t["bas"])


def ucret(gun: int, dl: list[dict]) -> tuple[Decimal, int, Decimal]:
    """(toplam ücret, serbest gün sayısı, gün'üncü günün günlük ücreti)"""
    toplam, serbest, gunluk = Z, 0, Z
    for d in dl:
        if d["ucret"] == 0 and d["bas"] == serbest + 1:
            serbest = d["bit"] if d["bit"] is not None else 10**6
        son = min(gun, d["bit"]) if d["bit"] is not None else gun
        if son >= d["bas"]:
            toplam += (son - d["bas"] + 1) * d["ucret"]
        if d["bas"] <= gun and (d["bit"] is None or gun <= d["bit"]):
            gunluk = d["ucret"]
    return toplam, serbest, gunluk


def takip(konteynerler: list[dict], tarife: list[dict], rapor: date, ilk_gun_dahil: bool = True) -> dict:
    if not konteynerler:
        raise ValueError("Takip edilecek konteyner yok.")
    ek, tarifesiz = 1 if ilk_gun_dahil else 0, set()
    for k in konteynerler:
        k["kalemler"] = {}
        for tur, bas, bit in (("demuraj", k["tahliye"], k["cikis"]), ("detention", k["cikis"], k["iade"])):
            if bas is None:
                continue
            dl = dilimler(tarife, tur, k["hat"], k["tip"])
            acik = bit is None
            gun = max(((bit or rapor) - bas).days + ek, 0)
            if not dl:
                tarifesiz.add((tur, k["hat"], k["tip"]))
                k["kalemler"][tur] = {"gun": gun, "acik": acik, "ucret": None, "serbest": None, "kalan": None, "son_serbest": None, "gunluk": None, "para": ""}
                continue
            u, serbest, gunluk = ucret(gun, dl)
            k["kalemler"][tur] = {"gun": gun, "acik": acik, "ucret": u, "serbest": serbest, "kalan": max(serbest - gun, 0) if acik else None,
                                  "son_serbest": bas + timedelta(days=serbest - ek) if acik and serbest < 10**6 else None, "gunluk": gunluk if acik else None, "para": next((d["para"] for d in dl if d["para"]), "")}
        d, t = k["kalemler"].get("demuraj"), k["kalemler"].get("detention")
        k["toplam"] = sum((x["ucret"] for x in (d, t) if x and x["ucret"]), Z)
        acik = t if t and t["acik"] else d if d["acik"] else None
        k["acik"] = "Detention" if acik is t and acik else "Demuraj" if acik else ""
        k["durum"] = "Kapalı" if not acik else "Tarife yok" if acik["ucret"] is None else "Ücret işliyor" if acik["gunluk"] else "Serbest süre içinde"
        k["acik_kalem"] = acik
    sira = ("Ücret işliyor", "Serbest süre içinde", "Tarife yok", "Kapalı")
    para = {x["para"] for k in konteynerler for x in k["kalemler"].values() if x["ucret"]}
    top = lambda tur: sum((k["kalemler"][tur]["ucret"] for k in konteynerler if tur in k["kalemler"] and k["kalemler"][tur]["ucret"]), Z)   # noqa: E731
    return {"konteynerler": sorted(konteynerler, key=lambda k: (sira.index(k["durum"]), (k["acik_kalem"] or {}).get("kalan") or 0, -k["toplam"], k["no"])), "rapor": rapor, "ilk_gun_dahil": ilk_gun_dahil,
            "toplam": {"demuraj": top("demuraj"), "detention": top("detention"), "genel": top("demuraj") + top("detention"),
                       "gunluk": sum((k["acik_kalem"]["gunluk"] for k in konteynerler if k["acik_kalem"] and k["acik_kalem"]["gunluk"]), Z)},
            "say": {s: sum(k["durum"] == s for k in konteynerler) for s in sira}, "paralar": sorted(para), "tarifesiz": sorted(tarifesiz)}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    renk = {"Ücret işliyor": KIRMIZI, "Serbest süre içinde": SARI, "Tarife yok": MAVI}
    ws = wb.active
    ws.title = "Konteyner Takibi"
    _baslik(ws, ["Durum", "Konteyner No", "Tip", "Hat", "Dosya / Müşteri", "Tahliye", "Limandan Çıkış", "Boş İade", "Demuraj Günü", "Demuraj Ücreti", "Detention Günü", "Detention Ücreti", "Toplam Ücret",
                 "Para Birimi", "Açık Kalem", "Kalan Serbest Gün", "Son Serbest Gün", "Bugünkü Günlük Ücret"], (18, 16, 8, 14, 18, 12, 12, 12, 9, 13, 9, 13, 13, 8, 11, 10, 12, 12))
    for k in r["konteynerler"]:
        d, t, a = k["kalemler"].get("demuraj"), k["kalemler"].get("detention"), k["acik_kalem"] or {}
        _satir(ws, [k["durum"], k["no"], k["tip"], k["hat"], k["ref"], k["tahliye"], k["cikis"], k["iade"], d["gun"], d["ucret"], t["gun"] if t else None, t["ucret"] if t else None, k["toplam"],
                    (d["para"] or (t or {}).get("para", "")), k["acik"], a.get("kalan"), a.get("son_serbest"), a.get("gunluk") or None], para_s=(10, 12, 13, 18), tarih_s=(6, 7, 8, 17),
               dolgu={1: renk.get(k["durum"])})
    t = r["toplam"]
    _satir(ws, ["", "Toplam", "", "", "", None, None, None, None, t["demuraj"], None, t["detention"], t["genel"], "", "", None, None, t["gunluk"] or None], para_s=(10, 12, 13, 18))
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.append([])
    ws.append([f"Rapor tarihi {r['rapor']:%d.%m.%Y}. Gün sayımında başlangıç günü {'dahil' if r['ilk_gun_dahil'] else 'hariç'}. Açık konteynerlerde ücret rapor tarihine kadar işlemiş tutardır; "
               "her gün artar. Serbest süre ve ücretler tarife dosyanızdan alınmıştır; hattın faturası esastır."])
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, tarife_yol: Path, cikti: Path, rapor: date, ilk_gun_dahil: bool = True) -> dict:
    tarife, uy = tarife_oku(tarife_yol)
    konteynerler, uy2 = oku(yol)
    r = takip(konteynerler, tarife, rapor, ilk_gun_dahil)
    uy += uy2
    for tur, hat, tip in r["tarifesiz"]:
        uy.append(uyari("Yüksek", "Tarife bulunamadı", f"{hat or '(hat yok)'} · {tip or '(tip yok)'}", f"{tur.capitalize()} için bu hat ve tipe uyan tarife satırı yok; ücret hesaplanmadı."))
    if len(r["paralar"]) > 1:
        uy.append(uyari("Orta", "Birden çok para birimi", tarife_yol.name, f"{', '.join(p or '(boş)' for p in r['paralar'])}: toplam satırı farklı para birimlerini toplar; kur çevirisi yapılmaz."))
    for k in r["konteynerler"]:
        if k["tahliye"] > rapor:
            uy.append(uyari("Orta", "Tahliye tarihi ileri tarihli", k["no"], f"Tahliye {k['tahliye']:%d.%m.%Y}, rapor tarihi {rapor:%d.%m.%Y}."))
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Konteyner demuraj ve detention sürelerini ve ücretlerini hesaplar.")
    p.add_argument("--konteynerler", type=Path, default=ORNEK / "konteynerler.csv", help="Konteyner No, Tip, Hat, Dosya No, Tahliye Tarihi, Limandan Çıkış, Boş İade")
    p.add_argument("--tarife", type=Path, default=ORNEK / "tarife.csv", help="Hat, Tip, Ücret Türü (Demuraj / Detention), Gün Başlangıç, Gün Bitiş, Günlük Ücret, Para Birimi")
    p.add_argument("--gun-sayimi", choices=("dahil", "haric"), default="dahil", help="Başlangıç günü (tahliye / çıkış günü) sayılsın mı (varsayılan dahil)")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 12.10.2026)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "demuraj_detention.xlsx")
    a = p.parse_args(argv)
    rapor = tarih(a.tarih) if a.tarih else (date(2026, 10, 12) if a.konteynerler == ORNEK / "konteynerler.csv" else date.today())
    if rapor is None:
        print("[X] --tarih GG.AA.YYYY biçiminde olmalı")
        return 2
    for y in (a.konteynerler, a.tarife):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.konteynerler, a.tarife, a.cikti, rapor, a.gun_sayimi == "dahil")
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t, s = r["toplam"], r["say"]
    print(f"[OK] {len(r['konteynerler'])} konteyner · {s['Ücret işliyor']} ücret işliyor, {s['Serbest süre içinde']} serbest sürede, {s['Kapalı']} kapalı")
    print(f"[{'!' if t['gunluk'] else 'OK'}] Demuraj {t['demuraj']:,.2f} · detention {t['detention']:,.2f} · toplam {t['genel']:,.2f} · açık konteynerlerde günlük işleyen {t['gunluk']:,.2f}")
    for k in r["konteynerler"]:
        a2 = k["acik_kalem"]
        if k["durum"] == "Ücret işliyor":
            print(f"[!] {k['no']}: {k['acik']} {a2['gun']}. gün, işlemiş {a2['ucret']:,.2f}, günlük {a2['gunluk']:,.2f}")
        elif k["durum"] == "Serbest süre içinde":
            print(f"     {k['no']}: {k['acik']} serbest sürede, {a2['kalan']} gün kaldı" + (f" (son serbest gün {a2['son_serbest']:%d.%m.%Y})" if a2["son_serbest"] else ""))
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
