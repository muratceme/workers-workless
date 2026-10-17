"""
Kanal Maliyet Analizi — Gerisi Bende kod bloğu
Rezervasyon ve Gelir Yönetimi › Gelir Müdürü (Revenue Manager)
OTA, acente ve direkt kanallardan gelen oda gelirini kanal maliyetleri düşüldükten sonra karşılaştırır.
  Komisyon        = tutar × komisyon % (satırda Komisyon % yazılıysa o, yoksa kanal tanımındaki oran)
  İşlem ücreti    = tutar × işlem ücreti % (sanal POS, ödeme altyapısı vb.)
  Rezervasyon ücreti = rezervasyon sayısı × rezervasyon başı ücret (GDS, kanal yöneticisi vb.)
  Sabit maliyet   = aylık sabit maliyet × dönemdeki ay sayısı (giriş tarihlerinin kapsadığı ay ya da --ay)
  Net gelir       = brüt gelir − toplam kanal maliyeti;  maliyet % = toplam maliyet ÷ brüt × 100
  ADR             = brüt gelir ÷ oda-gece;  net ADR = net gelir ÷ oda-gece
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
REZ = {"no": ("rezervasyon no", "rez no", "voucher"), "kanal": ("kanal", "satis kanali", "acente", "kaynak"), "giris": ("giris tarihi", "giris", "check in"),
       "cikis": ("cikis tarihi", "cikis", "check out"), "gece": ("gece", "gece sayisi", "oda gece"), "tutar": ("tutar", "oda geliri", "brut tutar", "gelir", "toplam tutar"),
       "komisyon": ("komisyon", "komisyon orani", "komisyon yuzde"), "durum": ("durum", "statu")}
KANAL = {"kanal": ("kanal", "kanal adi"), "tur": ("kanal turu", "tur", "tip"), "komisyon": ("komisyon", "komisyon orani", "komisyon yuzde"),
         "islem": ("islem ucreti", "islem ucreti yuzde", "odeme ucreti", "pos"), "rez_ucret": ("rezervasyon basi ucret", "rezervasyon ucreti", "rez basi ucret"),
         "sabit": ("aylik sabit maliyet", "sabit maliyet", "aylik ucret")}
SIFIR = Decimal(0)


def _yuzde(a: Decimal, b: Decimal) -> Decimal | None:
    return yuvarla(a * 100 / b) if b else None


def kanal_oku(yol: Path) -> tuple[dict, list[dict]]:
    uy, k = [], {}
    for r in kayitlar(yol, KANAL, ("kanal",)):
        ad = metin(r["kanal"])
        if not ad:
            continue
        d = {a: (para(r[a]) if metin(r[a]) else SIFIR) for a in ("komisyon", "islem", "rez_ucret", "sabit")}
        if any(v is None or v < 0 for v in d.values()) or d["komisyon"] + d["islem"] >= 100:
            uy.append(uyari("Yüksek", "Kanal satırı okunamadı", ad, f"Satır {r['_satir']}: oranlardan ya da ücretlerden biri okunamadı / geçersiz."))
            continue
        if katla(ad) in k:
            uy.append(uyari("Orta", "Kanal iki kez tanımlı", ad, "İkinci satır alınmadı."))
            continue
        k[katla(ad)] = {"kanal": ad, "tur": metin(r["tur"]), **d}
    return k, uy


def rez_oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, REZ, ("kanal", "tutar")):
        if not metin(r["kanal"]) and not metin(r["tutar"]):
            continue
        kim = metin(r["no"]) or f"Satır {r['_satir']}"
        d = katla(r["durum"])
        if "iptal" in d or "no show" in d or d in ("noshow", "cancelled", "canceled"):
            continue
        g, c = tarih(r["giris"]) if metin(r["giris"]) else None, tarih(r["cikis"]) if metin(r["cikis"]) else None
        tutar, gece = para(r["tutar"]), para(r["gece"]) if metin(r["gece"]) else None
        kom = para(r["komisyon"]) if metin(r["komisyon"]) else None
        if tutar is None or not metin(r["kanal"]) or (metin(r["giris"]) and g is None) or (metin(r["cikis"]) and c is None) or (metin(r["gece"]) and gece is None) or (metin(r["komisyon"]) and kom is None):
            uy.append(uyari("Yüksek", "Satır okunamadı", kim, f"Satır {r['_satir']}: kanal, tutar, tarih, gece ya da komisyon okunamadı. Rezervasyon alınmadı."))
            continue
        if metin(r["no"]):
            if katla(r["no"]) in gorulen:
                uy.append(uyari("Orta", "Rezervasyon iki kez yazılmış", kim, f"Satır {r['_satir']}: ikinci satır alınmadı."))
                continue
            gorulen.add(katla(r["no"]))
        if gece is None:
            gece = Decimal((c - g).days) if g and c and c > g else Decimal(1)
        l.append({"no": metin(r["no"]), "kanal": metin(r["kanal"]), "giris": g, "gece": gece, "tutar": tutar, "komisyon": kom})
    return l, uy


def analiz(rez: list[dict], kanallar: dict, ay: int | None = None, esik: Decimal = Decimal(25)) -> dict:
    if not rez:
        raise ValueError("Okunabilen rezervasyon yok.")
    uy = []
    aylar = {(x["giris"].year, x["giris"].month) for x in rez if x["giris"]}
    ay = ay or len(aylar) or 1
    gruplar = {}
    for x in rez:
        gruplar.setdefault(katla(x["kanal"]), []).append(x)
    satirlar = []
    for kk, l in gruplar.items():
        tanim = kanallar.get(kk)
        if tanim is None:
            uy.append(uyari("Orta", "Kanal tanımı yok", l[0]["kanal"], f"{len(l)} rezervasyon; kanal maliyet tablosunda yok. Satırda yazan komisyon dışında maliyet sıfır sayıldı."))
            tanim = {"kanal": l[0]["kanal"], "tur": "", "komisyon": SIFIR, "islem": SIFIR, "rez_ucret": SIFIR, "sabit": SIFIR}
        brut = sum((x["tutar"] for x in l), SIFIR)
        gece = sum((x["gece"] for x in l), SIFIR)
        kom = sum((yuvarla(x["tutar"] * (x["komisyon"] if x["komisyon"] is not None else tanim["komisyon"]) / 100) for x in l), SIFIR)
        farkli = sum(1 for x in l if x["komisyon"] is not None and abs(x["komisyon"] - tanim["komisyon"]) > 1)
        if farkli and kk in kanallar:
            uy.append(uyari("Bilgi", "Komisyon farklı", tanim["kanal"], f"{farkli} rezervasyonda satırdaki komisyon oranı kanal tanımından (%{tanim['komisyon']}) 1 puandan fazla farklı; satırdaki oran kullanıldı."))
        islem = yuvarla(brut * tanim["islem"] / 100)
        rez_ucret = yuvarla(tanim["rez_ucret"] * len(l))
        sabit = yuvarla(tanim["sabit"] * ay)
        maliyet = kom + islem + rez_ucret + sabit
        net = brut - maliyet
        satirlar.append({"kanal": tanim["kanal"], "tur": tanim["tur"], "rez": len(l), "gece": gece, "brut": brut, "adr": yuvarla(brut / gece) if gece else None, "komisyon": kom,
                         "islem": islem, "rez_ucret": rez_ucret, "sabit": sabit, "maliyet": maliyet, "maliyet_oran": _yuzde(maliyet, brut), "net": net,
                         "net_adr": yuvarla(net / gece) if gece else None, "rez_basi": yuvarla(maliyet / len(l))})
    for kk, t in kanallar.items():
        if kk not in gruplar and t["sabit"]:
            sabit = yuvarla(t["sabit"] * ay)
            satirlar.append({"kanal": t["kanal"], "tur": t["tur"], "rez": 0, "gece": SIFIR, "brut": SIFIR, "adr": None, "komisyon": SIFIR, "islem": SIFIR, "rez_ucret": SIFIR,
                             "sabit": sabit, "maliyet": sabit, "maliyet_oran": None, "net": -sabit, "net_adr": None, "rez_basi": None})
            uy.append(uyari("Orta", "Kanaldan satış yok", t["kanal"], f"Dönemde rezervasyon yok ama {sabit:,.2f} sabit maliyet var."))
    toplam_brut = sum((s["brut"] for s in satirlar), SIFIR)
    toplam_net = sum((s["net"] for s in satirlar), SIFIR)
    for s in satirlar:
        s["pay"] = _yuzde(s["brut"], toplam_brut)
        s["net_pay"] = _yuzde(s["net"], toplam_net)
        if s["maliyet_oran"] is not None and s["maliyet_oran"] > esik:
            uy.append(uyari("Bilgi", "Kanal maliyeti yüksek", s["kanal"], f"Kanal maliyeti brüt gelirin %{s['maliyet_oran']} kadarı (eşik %{esik}); net ADR {s['net_adr']:,.2f}."))
    satirlar.sort(key=lambda s: (-s["brut"], katla(s["kanal"])))
    sat = [s for s in satirlar if s["net_adr"] is not None]
    en_iyi = max(sat, key=lambda s: s["net_adr"]) if sat else None
    for s in satirlar:
        s["fark"] = (s["net_adr"] - en_iyi["net_adr"]) if en_iyi and s["net_adr"] is not None else None
    gece = sum((s["gece"] for s in satirlar), SIFIR)
    maliyet = sum((s["maliyet"] for s in satirlar), SIFIR)
    toplam = {"kanal": "TOPLAM", "tur": "", "rez": sum(s["rez"] for s in satirlar), "gece": gece, "brut": toplam_brut, "adr": yuvarla(toplam_brut / gece) if gece else None,
              **{a: sum((s[a] for s in satirlar), SIFIR) for a in ("komisyon", "islem", "rez_ucret", "sabit")}, "maliyet": maliyet, "maliyet_oran": _yuzde(maliyet, toplam_brut),
              "net": toplam_net, "net_adr": yuvarla(toplam_net / gece) if gece else None, "pay": Decimal(100) if toplam_brut else None, "net_pay": Decimal(100) if toplam_net else None, "fark": None}
    toplam["rez_basi"] = yuvarla(maliyet / toplam["rez"]) if toplam["rez"] else None
    return {"kanallar": satirlar, "toplam": toplam, "ay": ay, "en_iyi": en_iyi["kanal"] if en_iyi else None, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, esik: Decimal) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kanal Karşılaştırma"
    _baslik(ws, ["Kanal", "Tür", "Rezervasyon", "Oda-Gece", "Brüt Gelir", "Gelir Payı %", "ADR", "Komisyon", "İşlem Ücreti", "Rez. Başı Ücret", "Sabit Maliyet", "Toplam Maliyet",
                 "Maliyet %", "Net Gelir", "Net Payı %", "Net ADR", f"Net ADR Farkı ({r['en_iyi'] or '-'})", "Rezervasyon Başı Maliyet"],
            (20, 10, 11, 10, 14, 9, 10, 12, 11, 11, 11, 13, 9, 14, 9, 10, 13, 12), "B2")
    for s in r["kanallar"] + [r["toplam"]]:
        n = _satir(ws, [s["kanal"], s["tur"], s["rez"], s["gece"], s["brut"], s["pay"], s["adr"], s["komisyon"], s["islem"], s["rez_ucret"], s["sabit"], s["maliyet"], s["maliyet_oran"],
                        s["net"], s["net_pay"], s["net_adr"], s["fark"], s["rez_basi"]], para_s=tuple(range(5, 19)),
                   dolgu={13: KIRMIZI if s["maliyet_oran"] is not None and s["maliyet_oran"] > esik and s["kanal"] != "TOPLAM" else None})
        if s["kanal"] == "TOPLAM":
            for h in ws[n]:
                h.font = Font(bold=True)
    ws.append([])
    ws.append([f"Sabit maliyet {r['ay']} ay üzerinden. Net ADR farkı: en yüksek net ADR'li kanala göre oda-gece başına fark."])
    ws.append(["Tutarlar dosyadaki esasla (KDV dahil ya da hariç) hesaplanır; tüm kanallarda aynı esası kullanın."])
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(rez_yol: Path, kanal_yol: Path | None, cikti: Path, ay: int | None = None, esik: Decimal = Decimal(25)) -> dict:
    kanallar, uy1 = kanal_oku(kanal_yol) if kanal_yol else ({}, [])
    rez, uy2 = rez_oku(rez_yol)
    r = analiz(rez, kanallar, ay, esik)
    r["uyarilar"] = sirala(uy1 + uy2 + r["uyarilar"])
    rapor_yaz(cikti, r, esik)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Kanalların komisyon ve diğer maliyetler sonrası net getirisini karşılaştırır.")
    p.add_argument("--rezervasyonlar", type=Path, help="Rezervasyon No, Kanal, Giriş Tarihi, Çıkış Tarihi ya da Gece, Tutar, Komisyon %% (isteğe bağlı), Durum")
    p.add_argument("--kanallar", type=Path, help="Kanal, Kanal Türü, Komisyon %%, İşlem Ücreti %%, Rezervasyon Başı Ücret, Aylık Sabit Maliyet")
    p.add_argument("--ay", type=int, help="Sabit maliyet için ay sayısı (varsayılan: giriş tarihlerinin kapsadığı ay sayısı)")
    p.add_argument("--esik", default="25", help="Maliyet yüzdesi bu değerin üstündeyse bilgi notu (varsayılan 25)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kanal_maliyet_analizi.xlsx")
    a = p.parse_args(argv)
    rez = a.rezervasyonlar or ORNEK / "rezervasyonlar.csv"
    kanal = a.kanallar or (ORNEK / "kanallar.csv" if a.rezervasyonlar is None else None)
    esik = para(a.esik)
    if esik is None or not 0 <= esik <= 100 or (a.ay is not None and a.ay < 1):
        print("[X] --esik 0-100, --ay en az 1 olmalı")
        return 2
    for y in (rez, kanal):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(rez, kanal, a.cikti, a.ay, esik)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {t['rez']} rezervasyon · brüt {t['brut']:,.2f} · kanal maliyeti {t['maliyet']:,.2f} (%{t['maliyet_oran']}) · net {t['net']:,.2f}")
    for s in r["kanallar"]:
        print(f"     {s['kanal']}: brüt {s['brut']:,.2f} · maliyet " + (f"%{s['maliyet_oran']} · net ADR {s['net_adr']}" if s["rez"] else f"{s['maliyet']:,.2f} (satış yok)"))
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
