import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def r(kanal, ay, gece, tutar, kom=None, no=""):
    return {"no": no, "kanal": kanal, "giris": date(2026, ay, 5), "gece": D(gece), "tutar": D(tutar), "komisyon": None if kom is None else D(kom)}


def k(ad, kom=0, islem=0, rez=0, sabit=0):
    return {main.katla(ad): {"kanal": ad, "tur": "", "komisyon": D(kom), "islem": D(islem), "rez_ucret": D(rez), "sabit": D(sabit)}}


class AnalizTesti(unittest.TestCase):
    def test_hesap(self):
        kn = k("OTA", 18) | k("Web", 0, D("2.5"), 0, 1000) | k("GDS", 10, 0, 50) | k("Reklam", sabit=500)
        rez = [r("OTA", 9, 2, 6000), r("OTA", 10, 1, 3000, 20), r("Web", 9, 3, 9000), r("GDS", 10, 2, 8000), r("GDS", 10, 1, 4000), r("Ajans", 9, 2, 5000)]
        x = main.analiz(rez, kn)
        s = {a["kanal"]: a for a in x["kanallar"]}
        self.assertEqual(x["ay"], 2)
        # OTA: 6000×%18 + 3000×%20 = 1080 + 600
        self.assertEqual((s["OTA"]["komisyon"], s["OTA"]["maliyet_oran"], s["OTA"]["net"], s["OTA"]["adr"], s["OTA"]["net_adr"]), (D(1680), D("18.67"), D(7320), D(3000), D(2440)))
        # Web: işlem 225 + sabit 2 × 1000
        self.assertEqual((s["Web"]["islem"], s["Web"]["sabit"], s["Web"]["maliyet"], s["Web"]["net_adr"], s["Web"]["fark"]), (D(225), D(2000), D(2225), D("2258.33"), D("-1308.34")))
        self.assertEqual((s["GDS"]["komisyon"], s["GDS"]["rez_ucret"], s["GDS"]["maliyet"], s["GDS"]["rez_basi"], s["GDS"]["net_adr"]), (D(1200), D(100), D(1300), D(650), D("3566.67")))
        self.assertEqual((s["Reklam"]["rez"], s["Reklam"]["net"], s["Ajans"]["maliyet"]), (0, D(-1000), D(0)))
        self.assertEqual(x["en_iyi"], "GDS")
        t = x["toplam"]
        self.assertEqual((t["rez"], t["brut"], t["maliyet"], t["net"], t["maliyet_oran"]), (6, D(35000), D(6205), D(28795), D("17.73")))
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in x["uyarilar"]),
                         [("Bilgi", "Komisyon farklı", "OTA"), ("Orta", "Kanal tanımı yok", "Ajans"), ("Orta", "Kanaldan satış yok", "Reklam")])
        self.assertEqual(main.analiz(rez, kn, ay=1)["kanallar"][-1]["sabit"], D(500))
        with self.assertRaises(ValueError):
            main.analiz([], kn)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            x = main.calistir(ORNEK / "rezervasyonlar.csv", ORNEK / "kanallar.csv", cikti)
            t = x["toplam"]
            self.assertEqual((t["rez"], t["brut"], t["maliyet"], t["maliyet_oran"], x["ay"], x["en_iyi"]), (295, D("2632072.85"), D("414519.96"), D("15.75"), 2, "Telefon / e-posta"))
            self.assertEqual([s["kanal"] for s in x["kanallar"]][:3], ["OTA A", "Tur operatörü", "OTA B"])
            self.assertEqual(sorted(u["kim"] for u in x["uyarilar"] if u["onem"] == "Orta"), ["Etkinlik ajansı", "Sosyal medya reklam"])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Kanal Karşılaştırma", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            ky, ry = Path(tmp) / "k.csv", Path(tmp) / "r.csv"
            ky.write_text("Kanal;Komisyon Oranı;POS;Sabit Maliyet\nA;15;;1.000\nA;10;;\nB;x;;\nC;60;45;\n", encoding="utf-8")
            ry.write_text("Voucher;Satış Kanalı;Check In;Check Out;Oda Geliri;Durum\n1;A;01.09.2026;03.09.2026;1.000;\n1;A;01.09.2026;03.09.2026;1.000;\n2;A;;;500;İptal\n3;;;;;\n4;B;x;;100;\n5;B;;;200;\n", encoding="utf-8")
            kn, uy = main.kanal_oku(ky)
            self.assertEqual((list(kn), [u["tur"] for u in uy]), (["a"], ["Kanal iki kez tanımlı", "Kanal satırı okunamadı", "Kanal satırı okunamadı"]))
            l, uy = main.rez_oku(ry)
            self.assertEqual(([(x["no"], x["gece"], x["tutar"]) for x in l], [u["tur"] for u in uy]), ([("1", D(2), D(1000)), ("5", D(1), D(200))], ["Rezervasyon iki kez yazılmış", "Satır okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--rezervasyonlar", str(ry), "--kanallar", str(ky), "--cikti", c]
            self.assertEqual(main.main(arg + ["--ay", "3", "--esik", "10"]), 0)
            self.assertEqual(main.main(["--rezervasyonlar", str(ry), "--cikti", c]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--esik", "x"], ["--ay", "0"])], [2, 2])
            self.assertEqual(main.main(["--rezervasyonlar", str(Path(tmp) / "yok.csv")]), 1)
            ry.write_text("Voucher;Satış Kanalı;Oda Geliri;Durum\n1;A;5;İptal\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
