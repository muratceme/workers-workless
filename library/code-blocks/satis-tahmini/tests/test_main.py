import math
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
MEVSIM = [0.5, 0.6, 0.8, 1.0, 1.3, 1.6, 1.8, 1.6, 1.1, 0.8, 0.5, 0.4]


class YontemTesti(unittest.TestCase):
    def test_basit_yontemler(self):
        y = [10, 20, 30, 40, 50]
        self.assertEqual(main.naif(y, 2), [50, 50])
        self.assertEqual(main.ho3(y, 1), [40])
        self.assertEqual(main.mevsimsel_naif(list(range(1, 25)), 3), [13, 14, 15])
        t, _ = main.holt([float(5 + 3 * i) for i in range(20)], 3)      # tam doğrusal seri
        for a, b in zip(t, [65, 68, 71]):
            self.assertAlmostEqual(a, b, places=6)

    def test_holt_winters_mevsimsel_seri(self):
        y = [100 * (1 + 0.01 * i) * MEVSIM[i % 12] for i in range(48)]
        gercek = [100 * (1 + 0.01 * i) * MEVSIM[i % 12] for i in range(48, 54)]
        t, _ = main.holt_winters(y, 6, carpimsal=True)
        self.assertLess(main.wape(gercek, t), 0.03)
        self.assertGreater(main.wape(gercek, main.naif(y, 6)), 0.3)
        self.assertIsNone(main.holt_winters(y[:20], 6, carpimsal=False))       # < 24 ay
        self.assertIsNone(main.holt_winters([0.0] + y[1:], 6, carpimsal=True))  # çarpımsalda sıfır olamaz

    def test_yardimcilar(self):
        self.assertAlmostEqual(main.wape([100, 100], [90, 120]), 0.15)
        self.assertEqual(main.ay_coz("2026-09"), (2026, 9))
        self.assertEqual(main.ay_coz("09.2026"), (2026, 9))
        self.assertEqual(main.ay_coz("Eylül 2026"), (2026, 9))
        self.assertEqual(main.ay_coz("01.09.2026"), (2026, 9))
        self.assertEqual(main.ay_ileri((2026, 11), 3), (2027, 2))


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "s.xlsx"
        cls.s = main.calistir(ORNEK / "satislar.csv", cls.cikti)
        cls.x = {x.ad: x for x in cls.s["seriler"]}
        cls.u = {(u["tur"], u["kim"]) for u in cls.s["uyarilar"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_secim_ve_tutarlilik(self):
        for x in self.s["seriler"]:
            gecerli = {k: v["wape"] for k, v in x.sonuclar.items() if v["wape"] is not None}
            if gecerli:
                self.assertEqual(x.sonuclar[x.secilen]["wape"], min(gecerli.values()))
            self.assertEqual(len(x.tahmin), 12)
            self.assertTrue(all(v >= 0 for v in x.tahmin))
        self.assertAlmostEqual(sum(self.s["gelecek_toplam"]), sum(sum(x.tahmin) for x in self.s["seriler"]))
        self.assertEqual(self.s["gelecek"][0], (2026, 10))
        k = self.x["Klima · Marmara"]
        self.assertGreater(k.tahmin[9], k.tahmin[3])                 # Temmuz 2027 > Ocak 2027 (yaz zirvesi)

    def test_uyarilar(self):
        self.assertIn(("Aykırı ay", "Klima · Ege"), self.u)
        self.assertEqual(sum(1 for t, k in self.u if t == "Aykırı ay"), 1)    # yankı (Tem 2026) işaretlenmez
        self.assertIn(("Mevsimsellik yok", "Küçük ev aletleri · Akdeniz"), self.u)
        self.assertNotIn("hw_carpimsal", self.x["Küçük ev aletleri · Akdeniz"].sonuclar)
        m = self.s["mevsim"]["Klima"]
        self.assertGreater(m[6], 150)                                    # Temmuz endeksi
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Tahmin", "Yöntem Karşılaştırma", "Toplam", "Mevsimsellik", "Geçmiş + Tahmin", "Uyarılar"])

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx"), "--deger", "Tutar", "--ufuk", "6"]), 0)
            self.assertEqual(main.main(["--ufuk", "0"]), 2)
            self.assertEqual(main.main(["--satislar", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
