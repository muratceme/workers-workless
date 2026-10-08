import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 8)


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "f.xlsx"
        cls.s = main.calistir(ORNEK / "fason_isler.csv", ORNEK / "donusler.csv", cls.cikti, BUGUN)
        cls.i = {i.no: i for i in cls.s["isler"]}
        cls.u = {}
        for u in cls.s["uyarilar"]:
            cls.u.setdefault(u["tur"], []).append(u)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_durumlar(self):
        self.assertEqual({k: (i.durum, i.acik, i.gecikme) for k, i in self.i.items()},
                         {"FS-01": ("Geç tamamlandı", 0, 2), "FS-02": ("Devam ediyor", 2800, None), "FS-03": ("GECİKTİ", 1200, 8),
                          "FS-04": ("Tamamlandı", 0, -1), "FS-05": ("Tamamlandı", 0, -1), "FS-06": ("GECİKTİ", 2500, 1),
                          "FS-07": ("Yaklaşıyor", 900, None), "FS-08": ("Devam ediyor", 2000, None), "FS-09": ("Tamamlandı", 0, 0),
                          "FS-10": ("Eksik kapandı", 0, 2)})

    def test_uyarilar(self):
        hata = {u["is"]: u for u in self.u["Hata oranı yüksek"]}
        self.assertEqual(sorted(hata), ["FS-03", "FS-07"])
        self.assertIn("130 / 1800 hatalı (%7,2, eşik %3)", hata["FS-03"]["aciklama"])
        self.assertEqual(sorted(u["is"] for u in self.u["Tempo düşük"]), ["FS-02", "FS-07"])
        self.assertIn("100 adet dönmedi (%3,3", self.u["Eksik dönüş"][0]["aciklama"])
        self.assertEqual(self.u["Fazla dönüş"][0]["is"], "FS-09")
        self.assertEqual(self.u["Tamir bekliyor"][0]["is"], "FS-07")
        self.assertEqual(self.u["Tolerans içi fire"][0]["is"], "FS-01")
        self.assertIn("hiç dönüş yok", next(u for u in self.u["Gecikti"] if u["is"] == "FS-06")["aciklama"])

    def test_performans_ve_hakedis(self):
        p = {x["ad"]: x for x in self.s["performans"]}
        self.assertEqual((p["Atölye B"]["acik_adet"], p["Atölye B"]["kayip"], p["Atölye B"]["geciken_acik"]), (2100, 100, 1))
        self.assertAlmostEqual(p["Atölye B"]["hata_orani"], 265 / 5300)
        fs01 = sum(h["tutar"] for h in self.s["hakedis"] if h["is"].no == "FS-01")
        self.assertEqual(fs01, D(3910) * 18)
        self.assertEqual(sum(h["tutar"] for h in self.s["hakedis"]), D("447464"))

    def test_excel(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Açık İşler", "Tüm İşler", "Dönüşler", "Atölye Performansı", "Hakediş", "Uyarılar"])
        self.assertEqual([wb["Açık İşler"].cell(r, 1).value for r in range(2, 7)], ["FS-03", "FS-06", "FS-07", "FS-02", "FS-08"])
        self.assertEqual(wb["Açık İşler"]["P1"].value, "Atölyeyle Görüşme / Not")


class KuralTesti(unittest.TestCase):
    def test_donem_ve_veri_hatalari(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t / "i.csv").write_text("İş No;Atölye;Çıkış Tarihi;Çıkan Adet;Termin;Birim Fiyat (TL)\nA;X;10.09.2026;100;30.09.2026;10\nB;Y;10.09.2026;50;;\n",
                                     encoding="utf-8")
            (t / "d.csv").write_text("İş No;Dönüş Tarihi;Dönen Adet;Hatalı Adet\nA;05.09.2026;40;0\nA;02.10.2026;60;70\nZ;01.10.2026;5;0\n", encoding="utf-8")
            s = main.calistir(t / "i.csv", t / "d.csv", t / "o.xlsx", BUGUN, donem=(2026, 10))
            turler = {(u["tur"], u["is"]) for u in s["uyarilar"]}
            for b in (("Veri hatası", "A"), ("Tanımsız iş", "Z"), ("Termin yok", "B")):
                self.assertIn(b, turler)
            self.assertEqual([(h["donus"].adet, h["saglam"]) for h in s["hakedis"]], [(60, 0)])

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--isler", str(Path(tmp) / "yok.csv")]), 1)
            self.assertEqual(main.main(["--donem", "2026-13", "--cikti", str(Path(tmp) / "x.xlsx")]), 1)


if __name__ == "__main__":
    unittest.main()
