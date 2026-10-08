"""
Fason Atölye İş Takibi — Workers / Workless kod bloğu
Tekstil ve Konfeksiyon › Üretim ve Fason Takip › Fason Takip Sorumlusu

Fason atölyelere (dikim, baskı, nakış, yıkama, ütü-paket) çıkan işleri adet, çıkış-dönüş tarihi ve kalite durumuyla
takip eder:
  - Açık adet = çıkan − dönen; termini geçmiş açık iş (GECİKTİ), termine 2 gün kalmış açık iş (Yaklaşıyor).
  - Tempo: sürenin en az yarısı geçtiği hâlde dönüş oranı geçen süre oranının 30 puan gerisinde.
  - Kalite: dönüşteki hatalı adet oranı eşiği (varsayılan %3) aşıyor; "Tamir" sonucu bekleyen adetler.
  - Kapanış: "Kapandı" işaretli işte eksik dönen adet fire toleransını (varsayılan %1) aşıyor; çıkandan fazla dönüş.
  - Hakediş: sağlam adet (dönen − hatalı) × birim fiyat, atölye bazında; isteğe bağlı dönem (ay) filtresi.
Rapor: açık işler, tüm işler, dönüşler, atölye performansı, hakediş, uyarılar. İnternete bağlanmaz.

Kullanım:
    python main.py                                                   # örnek: 5 atölye, 10 iş, durum tarihi 08.10.2026
    python main.py --isler fason_isler.xlsx --donusler donusler.xlsx --bugun 08.10.2026 --donem 2026-09
"""
from __future__ import annotations

import argparse
import csv
import re
import statistics
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")

IS_SUTUNLARI = {"no": ("is no", "fason no", "is emri no", "no"), "atolye": ("atolye", "fasoncu", "tedarikci"), "siparis": ("siparis no", "po"),
                "model": ("model", "style"), "islem": ("islem", "operasyon", "is turu"), "cikis": ("cikis tarihi", "cikis", "gonderim tarihi"),
                "adet": ("cikan adet", "adet", "gonderilen adet"), "termin": ("termin", "donus termini", "teslim termini"),
                "fiyat": ("birim fiyat tl", "birim fiyat", "fiyat"), "kapandi": ("kapandi", "kapali", "durum")}
DONUS_SUTUNLARI = {"no": ("is no", "fason no", "is emri no", "no"), "tarih": ("donus tarihi", "tarih", "teslim tarihi"),
                   "adet": ("donen adet", "adet", "gelen adet"), "kalite": ("kalite sonucu", "kalite", "sonuc"),
                   "hatali": ("hatali adet", "hatali", "red adet"), "aciklama": ("hata aciklamasi", "aciklama", "hata")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def sayi(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tl(x: Decimal) -> str:
    return f"{x.quantize(Decimal('0.01'), ROUND_HALF_UP):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def yuzde(x: float) -> str:
    return f"{x * 100:.1f}".replace(".", ",").removesuffix(",0")


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Donus:
    satir: int
    no: str
    tarih: date | None
    adet: int
    kalite: str
    hatali: int
    aciklama: str


@dataclass
class Is:
    satir: int
    no: str
    atolye: str
    siparis: str
    model: str
    islem: str
    cikis: date | None
    adet: int
    termin: date | None
    fiyat: Decimal | None
    kapandi: bool
    donusler: list = field(default_factory=list)
    durum: str = ""
    gecikme: int | None = None

    @property
    def donen(self) -> int:
        return sum(d.adet for d in self.donusler)

    @property
    def hatali(self) -> int:
        return sum(d.hatali for d in self.donusler)

    @property
    def tamir(self) -> int:
        return sum(d.hatali for d in self.donusler if d.kalite == "Tamir")

    @property
    def acik(self) -> int:
        return 0 if self.kapandi else max(self.adet - self.donen, 0)

    @property
    def son_donus(self) -> date | None:
        return max((d.tarih for d in self.donusler if d.tarih), default=None)

    @property
    def tamamlandi(self) -> bool:
        return self.kapandi or self.donen >= self.adet


def kalite_bul(x) -> str:
    k = katla(x)
    if k.startswith(("tamir", "rework", "duzelt")):
        return "Tamir"
    if k.startswith(("red", "ret", "reject")):
        return "Red"
    return "Kabul" if k else ""


def isleri_oku(yol: Path) -> list[Is]:
    sonuc = []
    for r in kayitlar(yol, IS_SUTUNLARI, ("no", "adet")):
        if not r.get("no"):
            continue
        kap = katla(r.get("kapandi"))
        sonuc.append(Is(r["_satir"], metin(r["no"]), metin(r.get("atolye")), metin(r.get("siparis")), metin(r.get("model")), metin(r.get("islem")),
                        tarih(r.get("cikis")), int(sayi(r.get("adet")) or 0), tarih(r.get("termin")), sayi(r.get("fiyat")),
                        kap in ("evet", "e", "x", "kapandi", "kapali", "tamam", "1", "yes")))
    return sonuc


def donusleri_oku(yol: Path | None) -> list[Donus]:
    if yol is None:
        return []
    return [Donus(r["_satir"], metin(r["no"]), tarih(r.get("tarih")), int(sayi(r.get("adet")) or 0), kalite_bul(r.get("kalite")),
                  int(sayi(r.get("hatali")) or 0), metin(r.get("aciklama")))
            for r in kayitlar(yol, DONUS_SUTUNLARI, ("no", "adet")) if r.get("no")]


# ----------------------------------------------------------------------------
# Analiz
# ----------------------------------------------------------------------------

SIRA = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}


def analiz_et(isler: list[Is], donusler: list[Donus], bugun: date, uyari_gun: int = 2, hata_esik: float = 0.03, fire_tolerans: float = 0.01,
              tempo_puan: float = 0.30, donem: tuple[int, int] | None = None) -> dict:
    uyarilar = []

    def uyar(onem, tur, i: Is | None, aciklama, no=""):
        uyarilar.append({"onem": onem, "tur": tur, "is": i.no if i else no, "atolye": i.atolye if i else "", "siparis": i.siparis if i else "",
                         "aciklama": aciklama})

    sozluk = {i.no: i for i in isler}
    for d in donusler:
        if d.no in sozluk:
            sozluk[d.no].donusler.append(d)
        else:
            uyar("Orta", "Tanımsız iş", None, f"Satır {d.satir}: {d.adet} adet dönüş var ama iş listesinde '{d.no}' yok", d.no)
        if d.hatali > d.adet:
            uyar("Orta", "Veri hatası", sozluk.get(d.no), f"Satır {d.satir}: hatalı adet ({d.hatali}) dönen adetten ({d.adet}) büyük", d.no)
    for i in isler:
        i.donusler.sort(key=lambda d: d.tarih or date.min)
        for d in i.donusler:
            if d.tarih and i.cikis and d.tarih < i.cikis:
                uyar("Orta", "Veri hatası", i, f"{d.tarih:%d.%m.%Y} dönüşü çıkış tarihinden ({i.cikis:%d.%m.%Y}) önce")
        if i.donen > i.adet:
            uyar("Orta", "Fazla dönüş", i, f"Çıkan {i.adet}, dönen {i.donen} (+{i.donen - i.adet}). Başka işle karışmış veya çıkış eksik yazılmış olabilir; "
                 "sayımı kontrol edin.")
        # Kalite
        if i.donen:
            oran = i.hatali / i.donen
            if oran > hata_esik:
                aciklamalar = ", ".join(dict.fromkeys(p.strip() for d in i.donusler if d.hatali for p in d.aciklama.split(",") if p.strip()))
                uyar("Yüksek" if oran >= 2 * hata_esik else "Orta", "Hata oranı yüksek", i, f"{i.hatali} / {i.donen} hatalı (%{yuzde(oran)}, eşik "
                     f"%{yuzde(hata_esik)}). Hatalar: {aciklamalar or '—'}")
        if i.tamir:
            uyar("Orta", "Tamir bekliyor", i, f"{i.tamir} adet tamire ayrıldı. Tamirden dönenleri yeni dönüş satırı olarak girin.")
        # Durum
        if i.tamamlandi:
            i.durum = "Tamamlandı"
            son = i.son_donus
            if i.termin and son:
                i.gecikme = (son - i.termin).days
                if i.gecikme > 0:
                    i.durum = "Geç tamamlandı"
            if i.kapandi and i.donen < i.adet:
                eksik = i.adet - i.donen
                if eksik / i.adet > fire_tolerans:
                    i.durum = "Eksik kapandı"
                    uyar("Yüksek", "Eksik dönüş", i, f"İş kapatıldı ama {eksik} adet dönmedi (%{yuzde(eksik / i.adet)}, tolerans %{yuzde(fire_tolerans)}). "
                         f"Kayıp adet için atölyeyle mutabakat yapın" + (f"; hakedişten mahsup edilecek tutarı belirleyin." if i.fiyat else "."))
                else:
                    uyar("Bilgi", "Tolerans içi fire", i, f"{eksik} adet dönmedi (%{yuzde(eksik / i.adet)}), tolerans içinde")
            continue
        if not i.termin:
            i.durum = "Açık (termin yok)"
            uyar("Bilgi", "Termin yok", i, "Dönüş termini yazılmamış")
            continue
        if bugun > i.termin:
            i.durum, i.gecikme = "GECİKTİ", (bugun - i.termin).days
            uyar("Yüksek", "Gecikti", i, f"Termin {i.termin:%d.%m.%Y}, {i.gecikme} gün geçti; {i.acik} / {i.adet} adet hâlâ atölyede"
                 + (f" (son dönüş {i.son_donus:%d.%m.%Y})" if i.son_donus else " (hiç dönüş yok)"))
            continue
        kalan = (i.termin - bugun).days
        if kalan <= uyari_gun:
            i.durum = "Yaklaşıyor"
            uyar("Orta", "Termin yaklaşıyor", i, f"Termin {i.termin:%d.%m.%Y} ({kalan} gün); {i.acik} / {i.adet} adet açık")
        else:
            i.durum = "Devam ediyor"
        if i.cikis and i.termin > i.cikis:
            sure_orani = (bugun - i.cikis).days / (i.termin - i.cikis).days
            ilerleme = i.donen / i.adet if i.adet else 1
            if sure_orani >= 0.5 and ilerleme < sure_orani - tempo_puan:
                uyar("Orta", "Tempo düşük", i, f"Geçen süre %{sure_orani * 100:.0f}, dönen adet %{ilerleme * 100:.0f}; bu tempoyla termin riskli. "
                     "Atölyeden günlük çıkış adedi isteyin.")

    # Atölye performansı
    perf = []
    gruplar = defaultdict(list)
    for i in isler:
        gruplar[i.atolye or "—"].append(i)
    for ad, lst in sorted(gruplar.items()):
        biten = [i for i in lst if i.tamamlandi and i.gecikme is not None]
        gec = [i.gecikme for i in biten if i.gecikme > 0]
        donen = sum(i.donen for i in lst)
        perf.append({"ad": ad, "is": len(lst), "acik_is": sum(not i.tamamlandi for i in lst), "acik_adet": sum(i.acik for i in lst),
                     "biten": len(biten), "zamaninda": len(biten) - len(gec), "ort_gecikme": statistics.fmean(gec) if gec else 0,
                     "geciken_acik": sum(i.durum == "GECİKTİ" for i in lst), "donen": donen,
                     "hata_orani": sum(i.hatali for i in lst) / donen if donen else None,
                     "kayip": sum(i.adet - i.donen for i in lst if i.kapandi and i.donen < i.adet)})

    # Hakediş
    hakedis = []
    for i in isler:
        for d in i.donusler:
            if donem and not (d.tarih and (d.tarih.year, d.tarih.month) == donem):
                continue
            saglam = max(d.adet - d.hatali, 0)
            hakedis.append({"is": i, "donus": d, "saglam": saglam, "tutar": saglam * i.fiyat if i.fiyat is not None else None})
            if i.fiyat is None:
                uyar("Orta", "Fiyat yok", i, "Birim fiyat boş; hakediş hesaplanamadı")
    uyarilar.sort(key=lambda u: (SIRA[u["onem"]], u["atolye"], u["is"]))
    return {"uyarilar": uyarilar, "performans": perf, "hakedis": hakedis}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE", "GECİKTİ": "FDE2E1", "Eksik kapandı": "FDE2E1", "Yaklaşıyor": "FFF4CE",
        "Geç tamamlandı": "FFF4CE", "Tamamlandı": "E3F4E1", "Devam ediyor": "E8F0FE"}
KONTROL = PatternFill("solid", fgColor="FFF4CE")
UST = Alignment(vertical="top", wrap_text=True)
TF = "DD.MM.YYYY"
PF = "#,##0.00"


def _baslik(ws, basliklar, genislik=()):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def _is_satiri(i: Is) -> list:
    return [i.no, i.atolye, i.siparis, i.model, i.islem, i.cikis, i.termin, i.adet, i.donen, i.acik, i.hatali,
            i.hatali / i.donen if i.donen else None, i.son_donus, i.gecikme, i.durum]


IS_BASLIK = ["İş No", "Atölye", "Sipariş", "Model", "İşlem", "Çıkış", "Termin", "Çıkan", "Dönen", "Açık", "Hatalı", "Hata %", "Son Dönüş",
             "Gecikme (gün)", "Durum"]
IS_GENISLIK = (8, 12, 9, 8, 9, 11, 11, 8, 8, 8, 7, 7, 11, 9, 14)


def _bicim(ws, i: Is):
    r = ws.max_row
    for j in (6, 7, 13):
        ws.cell(r, j).number_format = TF
    ws.cell(r, 12).number_format = "0.0%"
    if i.durum in RENK:
        ws.cell(r, 15).fill = PatternFill("solid", fgColor=RENK[i.durum])


def rapor_yaz(cikti: Path, isler: list[Is], s: dict, bugun: date, donem) -> None:
    wb = Workbook()
    ac = wb.active
    ac.title = "Açık İşler"
    _baslik(ac, IS_BASLIK + ["Atölyeyle Görüşme / Not"], IS_GENISLIK + (30,))
    acik = sorted((i for i in isler if not i.tamamlandi), key=lambda i: (i.termin or date.max))
    for i in acik:
        ac.append(_is_satiri(i) + [""])
        _bicim(ac, i)
        ac.cell(ac.max_row, 16).fill = KONTROL
    ac.append([])
    ac.append([f"Durum tarihi {bugun:%d.%m.%Y} · açık iş {len(acik)} · atölyede {sum(i.acik for i in acik):,} adet".replace(",", ".")])

    ti = wb.create_sheet("Tüm İşler")
    _baslik(ti, IS_BASLIK, IS_GENISLIK)
    for i in isler:
        ti.append(_is_satiri(i))
        _bicim(ti, i)
    ti.auto_filter.ref = f"A1:O{ti.max_row}"

    dn = wb.create_sheet("Dönüşler")
    _baslik(dn, ["İş No", "Atölye", "Model", "Dönüş Tarihi", "Dönen", "Kalite", "Hatalı", "Hata Açıklaması"], (8, 12, 8, 11, 8, 8, 7, 40))
    for i in isler:
        for d in i.donusler:
            dn.append([i.no, i.atolye, i.model, d.tarih, d.adet, d.kalite, d.hatali, d.aciklama])
            dn.cell(dn.max_row, 4).number_format = TF
    dn.auto_filter.ref = f"A1:H{dn.max_row}"

    ap = wb.create_sheet("Atölye Performansı")
    _baslik(ap, ["Atölye", "İş", "Açık İş", "Atölyedeki Adet", "Biten İş", "Zamanında Biten", "Zamanında %", "Geç Bitende Ort. Gecikme (gün)",
                 "Termini Geçmiş Açık İş", "Dönen Adet", "Hata %", "Kayıp Adet (kapanan işler)"], (14, 6, 8, 10, 8, 10, 10, 13, 12, 10, 8, 12))
    for x in s["performans"]:
        ap.append([x["ad"], x["is"], x["acik_is"], x["acik_adet"], x["biten"], x["zamaninda"], x["zamaninda"] / x["biten"] if x["biten"] else None,
                   round(x["ort_gecikme"], 1), x["geciken_acik"], x["donen"], x["hata_orani"], x["kayip"]])
        ap.cell(ap.max_row, 7).number_format = ap.cell(ap.max_row, 11).number_format = "0.0%"

    hk = wb.create_sheet("Hakediş")
    hk.append([f"Hakediş · {'dönem ' + f'{donem[1]:02d}.{donem[0]}' if donem else 'tüm dönüşler'} · sağlam adet = dönen − hatalı"])
    hk.cell(1, 1).font = Font(bold=True)
    _baslik(hk, ["Atölye", "İş No", "Model", "İşlem", "Dönüş Tarihi", "Dönen", "Hatalı", "Sağlam", "Birim Fiyat", "Tutar (TL)", "Onay"],
            (14, 8, 8, 9, 11, 8, 7, 8, 10, 14, 12))
    hk.freeze_panes = "A3"
    toplam = defaultdict(lambda: [0, Decimal(0)])
    for h in sorted(s["hakedis"], key=lambda h: (h["is"].atolye, h["is"].no, h["donus"].tarih or date.min)):
        i, d = h["is"], h["donus"]
        hk.append([i.atolye, i.no, i.model, i.islem, d.tarih, d.adet, d.hatali, h["saglam"], float(i.fiyat) if i.fiyat is not None else None,
                   float(h["tutar"]) if h["tutar"] is not None else None, ""])
        hk.cell(hk.max_row, 5).number_format = TF
        hk.cell(hk.max_row, 10).number_format = PF
        hk.cell(hk.max_row, 11).fill = KONTROL
        toplam[i.atolye][0] += h["saglam"]
        toplam[i.atolye][1] += h["tutar"] or 0
    hk.append([])
    hk.append(["Atölye toplamı", "", "", "", "", "", "", "Sağlam", "", "Tutar (TL)"])
    hk.cell(hk.max_row, 1).font = Font(bold=True)
    for ad, (adet, tutar) in sorted(toplam.items()):
        hk.append([ad, "", "", "", "", "", "", adet, "", float(tutar)])
        hk.cell(hk.max_row, 10).number_format = PF
        hk.cell(hk.max_row, 10).font = Font(bold=True)
    hk.append(["Not: KDV, stopaj ve kayıp adet mahsubu bu tabloya dahil değildir."])

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "İş No", "Atölye", "Sipariş", "Açıklama", "Aksiyon"], (9, 20, 8, 12, 9, 90, 24))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["is"], u["atolye"], u["siparis"], u["aciklama"], ""])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 6).alignment = UST
        uy.cell(uy.max_row, 7).fill = KONTROL
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(is_yolu: Path, donus_yolu: Path | None, cikti: Path, bugun: date | None = None, hata_esik: float = 0.03, fire_tolerans: float = 0.01,
             donem: tuple[int, int] | None = None, uyari_gun: int = 2) -> dict:
    bugun = bugun or date.today()
    isler = isleri_oku(is_yolu)
    donusler = donusleri_oku(donus_yolu)
    s = analiz_et(isler, donusler, bugun, uyari_gun, hata_esik, fire_tolerans, donem=donem)
    rapor_yaz(cikti, isler, s, bugun, donem)
    return {**s, "isler": isler, "bugun": bugun}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Fason atölyelere çıkan işleri adet, dönüş ve kalite durumuyla takip eder; gecikmeleri ve hakedişi listeler.")
    p.add_argument("--isler", type=Path, default=ORNEK / "fason_isler.csv", help="İş No, Atölye, Sipariş, Model, İşlem, Çıkış, Çıkan Adet, Termin, Birim Fiyat, Kapandı")
    p.add_argument("--donusler", type=Path, help="İş No, Dönüş Tarihi, Dönen Adet, Kalite Sonucu, Hatalı Adet, Hata Açıklaması")
    p.add_argument("--bugun", help="Durum tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 08.10.2026)")
    p.add_argument("--hata-esik", type=float, default=3, help="Hatalı adet oranı eşiği %% (varsayılan 3)")
    p.add_argument("--fire-tolerans", type=float, default=1, help="Kapanan işte kabul edilen eksik dönüş %% (varsayılan 1)")
    p.add_argument("--donem", help="Hakediş dönemi YYYY-AA (varsayılan tüm dönüşler)")
    p.add_argument("--uyari-gun", type=int, default=2, help="Termine kaç gün kala 'yaklaşıyor' denir (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "fason_takip.xlsx")
    a = p.parse_args(argv)
    ornek = a.isler == ORNEK / "fason_isler.csv"
    if ornek and a.donusler is None:
        a.donusler = ORNEK / "donusler.csv"
    for y in (a.isler, a.donusler):
        if y is not None and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    bugun = tarih(a.bugun) if a.bugun else (date(2026, 10, 8) if ornek else None)
    if a.bugun and bugun is None:
        print("[X] --bugun GG.AA.YYYY biçiminde olmalı")
        return 1
    donem = None
    if a.donem:
        m = re.fullmatch(r"(\d{4})-(\d{1,2})", a.donem.strip())
        if not m or not 1 <= int(m.group(2)) <= 12:
            print("[X] --donem YYYY-AA biçiminde olmalı, ör. 2026-09")
            return 1
        donem = (int(m.group(1)), int(m.group(2)))
    try:
        s = calistir(a.isler, a.donusler, a.cikti, bugun, a.hata_esik / 100, a.fire_tolerans / 100, donem, a.uyari_gun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    acik = [i for i in s["isler"] if not i.tamamlandi]
    print(f"[OK] {len(s['isler'])} iş · {len(acik)} açık · atölyede {sum(i.acik for i in acik)} adet · durum tarihi {s['bugun']:%d.%m.%Y}")
    for u in s["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['is']} {u['atolye']} {u['tur']}: {u['aciklama']}")
    toplam = sum((h["tutar"] for h in s["hakedis"] if h["tutar"] is not None), Decimal(0))
    print(f"[OK] Hakediş toplamı {tl(toplam)} TL · Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
