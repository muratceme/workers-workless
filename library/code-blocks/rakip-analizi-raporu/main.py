"""
Rakip Analizi Raporu — Workers / Workless kod bloğu
Pazarlama › Pazarlama Uzmanı

Rakip fiyat, ürün ve kampanya bilgilerini derleyip karşılaştırmalı rapor hazırlar:
  - Fiyat endeksi: rakibin birim fiyatı / bizim birim fiyatımız × 100 (ambalaj miktarına göre birim fiyat; indirimli
    fiyat varsa o). Her rakip-ürün için en son gözlem kullanılır. 100'ün üstü rakip daha pahalı demektir.
  - Ürün konumu: en ucuz / en pahalı / pazar ortası. Bizim birim fiyat rakiplerin medyanından --esik (%10) fazla
    yüksekse fiyat dezavantajı; tüm rakipler bizden %10'dan fazla pahalıysa "fiyat bırakma" bilgisi.
  - Rakip ve kategori özeti: medyan endeks, gözlem sayısı, stokta olmama oranı.
  - Fiyat değişimi: aynı rakip ürününün ilk ve son gözlemi.
  - Kampanyalar: rakip bazında sayı ve ortalama indirim, rapor tarihinde süren kampanyalar, bizim kategorimizde
    --derin (%30) ve üzeri indirim.
  - Özellik matrisi: bizde olmayıp rakiplerin çoğunda olan özellik (eksik) ve yalnız bizde olan özellik (avantaj).
Rapor: özet, fiyat karşılaştırma, rakip özeti, kategori özeti, fiyat değişimi, kampanyalar, özellik matrisi, veri,
uyarılar. İnternete bağlanmaz; veriler kullanıcı tarafından toplanır.

Kullanım:
    python main.py                                                 # örnek: 12 ürün, 4 rakip
    python main.py --urunler urunler.xlsx --fiyatlar rakip_fiyatlari.xlsx --kampanyalar kampanyalar.xlsx --ozellikler ozellikler.xlsx
"""
from __future__ import annotations

import argparse
import csv
import re
import statistics
import sys
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
K2 = Decimal("0.01")
BIZ = "Biz"

URUN_SUTUNLARI = {"kod": ("urun kodu", "kod", "sku"), "ad": ("urun", "urun adi"), "kategori": ("kategori",), "fiyat": ("fiyat", "satis fiyati", "liste fiyati"),
                  "miktar": ("ambalaj miktari", "miktar", "icerik"), "birim": ("birim", "olcu birimi")}
FIYAT_SUTUNLARI = {"tarih": ("tarih", "gozlem tarihi"), "rakip": ("rakip", "marka", "firma"), "kanal": ("kanal", "magaza", "site"),
                   "kod": ("bizim urun kodu", "eslesen urun kodu", "urun kodu"), "urun": ("rakip urun", "rakip urun adi", "urun"), "fiyat": ("fiyat", "liste fiyati"),
                   "indirimli": ("indirimli fiyat", "kampanyali fiyat"), "miktar": ("ambalaj miktari", "miktar", "icerik"), "birim": ("birim", "olcu birimi"),
                   "stok": ("stokta", "stok durumu", "stok")}
KAMPANYA_SUTUNLARI = {"bas": ("baslangic", "baslangic tarihi"), "bit": ("bitis", "bitis tarihi"), "rakip": ("rakip", "marka", "firma"), "kanal": ("kanal",),
                      "kampanya": ("kampanya", "aciklama"), "indirim": ("indirim", "indirim orani", "indirim yuzde"), "kategori": ("kategori",)}
OZELLIK_SUTUNLARI = {"marka": ("marka", "rakip", "firma"), "kategori": ("kategori", "urun grubu"), "ozellik": ("ozellik",), "deger": ("deger", "var yok")}
BIRIM_CARP = {"g": ("g", 1), "gr": ("g", 1), "kg": ("g", 1000), "ml": ("ml", 1), "cl": ("ml", 10), "l": ("ml", 1000), "lt": ("ml", 1000), "adet": ("adet", 1),
              "ad": ("adet", 1), "m": ("m", 1), "cm": ("m", Decimal("0.01")), "yikama": ("yikama", 1)}
VAR = {"var", "evet", "e", "x", "1", "true", "mevcut"}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("₺", "").replace("%", "").replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tl(x: Decimal | None) -> str:
    return "—" if x is None else f"{x.quantize(K2, ROUND_HALF_UP):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


def miktar_coz(miktar, birim) -> tuple[Decimal, str] | None:
    """('500', 'g') veya ('1,5 L', '') → (500, 'g') / (1500, 'ml')."""
    s = f"{metin(miktar)} {metin(birim)}".strip()
    m = re.match(r"^\s*(\d+(?:[.,]\d+)?)\s*([a-zA-ZçğıöşüÇĞİÖŞÜ]*)", s)
    if not m:
        return None
    sayi = Decimal(m[1].replace(",", "."))
    b = katla(m[2] or metin(birim))
    if b not in BIRIM_CARP:
        return None
    temel, carp = BIRIM_CARP[b]
    return sayi * carp, temel


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Urun:
    kod: str
    ad: str
    kategori: str
    fiyat: Decimal
    miktar: Decimal | None
    temel: str

    @property
    def birim_fiyat(self) -> Decimal:
        return self.fiyat / self.miktar if self.miktar else self.fiyat


@dataclass
class Gozlem:
    satir: int
    tarih: date
    rakip: str
    kanal: str
    kod: str
    urun: str
    fiyat: Decimal
    indirimli: Decimal | None
    miktar: Decimal | None
    temel: str
    stokta: bool | None

    @property
    def etkin(self) -> Decimal:
        return self.indirimli if self.indirimli is not None and self.indirimli < self.fiyat else self.fiyat


def urunleri_oku(yol: Path) -> tuple[dict[str, Urun], list[dict]]:
    urunler, uy = {}, []
    for r in kayitlar(yol, URUN_SUTUNLARI, ("kod", "fiyat")):
        kod, f = metin(r.get("kod")).upper(), para(r.get("fiyat"))
        if not kod or f is None:
            continue
        mk = miktar_coz(r.get("miktar"), r.get("birim")) if metin(r.get("miktar")) else None
        if metin(r.get("miktar")) and not mk:
            uy.append({"onem": "Orta", "tur": "Ambalaj miktarı okunamadı", "kim": kod, "aciklama": "Birim fiyat yerine paket fiyatı kullanıldı"})
        urunler[kod] = Urun(kod, metin(r.get("ad")), metin(r.get("kategori")) or "Diğer", f, mk[0] if mk else None, mk[1] if mk else "")
    return urunler, uy


def gozlemleri_oku(yol: Path, urunler: dict[str, Urun], bugun: date) -> tuple[list[Gozlem], list[dict]]:
    gozlemler, uy, eslesmeyen = [], [], defaultdict(int)
    for r in kayitlar(yol, FIYAT_SUTUNLARI, ("rakip", "fiyat")):
        rakip, f, t = metin(r.get("rakip")), para(r.get("fiyat")), tarih(r.get("tarih")) or bugun
        if not rakip or f is None:
            continue
        kod = metin(r.get("kod")).upper()
        if kod not in urunler:
            eslesmeyen[rakip] += 1
            continue
        mk = miktar_coz(r.get("miktar"), r.get("birim")) if metin(r.get("miktar")) else None
        s = katla(r.get("stok"))
        gozlemler.append(Gozlem(r["_satir"], t, rakip, metin(r.get("kanal")), kod, metin(r.get("urun")), f, para(r.get("indirimli")), mk[0] if mk else None,
                                mk[1] if mk else "", (s in VAR or s == "stokta") if s else None))
    for rk, n in eslesmeyen.items():
        uy.append({"onem": "Bilgi", "tur": "Eşleşmeyen fiyat", "kim": rk, "aciklama": f"{n} satırda 'Bizim Ürün Kodu' boş veya ürün listesinde yok; karşılaştırmaya alınmadı"})
    return gozlemler, uy


# ----------------------------------------------------------------------------
# Analiz
# ----------------------------------------------------------------------------

def endeks(g: Gozlem, u: Urun) -> tuple[Decimal | None, str]:
    """Rakip birim fiyatı / bizim birim fiyat × 100. Birimler uyuşmazsa paket fiyatı karşılaştırılır (not ile)."""
    if g.miktar and u.miktar and g.temel == u.temel:
        return (g.etkin / g.miktar) / u.birim_fiyat * 100, ""
    if g.miktar or u.miktar:
        return g.etkin / u.fiyat * 100, "birim uyuşmadı / eksik: paket fiyatı karşılaştırıldı"
    return g.etkin / u.fiyat * 100, ""


def analiz_et(urunler: dict[str, Urun], gozlemler: list[Gozlem], kampanyalar: list[dict], ozellikler: list[dict], bugun: date, esik: Decimal = Decimal("0.10"),
              derin: Decimal = Decimal("0.30")) -> dict:
    uy = []
    son = {}
    for g in sorted(gozlemler, key=lambda g: (g.tarih, g.satir)):
        son[(g.rakip, g.kod)] = g
    rakipler = sorted({g.rakip for g in gozlemler}, key=katla)
    karsilastirma = []
    for kod, u in sorted(urunler.items(), key=lambda i: (katla(i[1].kategori), i[0])):
        satir = {"urun": u, "rakip": {}, "notlar": []}
        for rk in rakipler:
            g = son.get((rk, kod))
            if g:
                e, n = endeks(g, u)
                satir["rakip"][rk] = {"g": g, "endeks": e}
                if n:
                    satir["notlar"].append(f"{rk}: {n}")
        endeksler = [v["endeks"] for v in satir["rakip"].values()]
        if endeksler:
            en_dusuk, en_yuksek = min(endeksler), max(endeksler)
            satir["medyan"] = statistics.median(endeksler)
            if en_dusuk > 100 and en_yuksek > 100:
                satir["konum"] = "En ucuz"
            elif en_dusuk < 100 and en_yuksek < 100:
                satir["konum"] = "En pahalı"
            else:
                satir["konum"] = "Pazar ortası"
            if 100 / satir["medyan"] - 1 > esik:
                ucuz = min(satir["rakip"].items(), key=lambda i: i[1]["endeks"])
                uy.append({"onem": "Orta", "tur": "Fiyat dezavantajı", "kim": f"{u.kod} {u.ad}",
                           "aciklama": f"Bizim birim fiyat rakiplerin medyanından %{(100 / satir['medyan'] - 1) * 100:.0f} yüksek "
                                       f"(medyan endeks {satir['medyan']:.0f}; en ucuz {ucuz[0]} {ucuz[1]['endeks']:.0f})"})
            if en_dusuk / 100 - 1 > esik:
                uy.append({"onem": "Bilgi", "tur": "Fiyat bırakma olasılığı", "kim": f"{u.kod} {u.ad}",
                           "aciklama": f"Tüm rakipler %{(en_dusuk / 100 - 1) * 100:.0f} veya daha fazla pahalı; fiyat artışı alanı olabilir"})
        else:
            satir["medyan"], satir["konum"] = None, "Rakip gözlemi yok"
        karsilastirma.append(satir)

    rakip_ozet = []
    for rk in rakipler:
        e = [s["rakip"][rk]["endeks"] for s in karsilastirma if rk in s["rakip"]]
        gl = [g for g in gozlemler if g.rakip == rk]
        stok = [g for g in gl if g.stokta is not None]
        rakip_ozet.append({"rakip": rk, "urun": len(e), "medyan": statistics.median(e) if e else None, "ucuz": sum(x < 100 for x in e), "pahali": sum(x > 100 for x in e),
                           "gozlem": len(gl), "stoksuz": sum(not g.stokta for g in stok) / len(stok) if stok else None,
                           "indirimli": sum(g.indirimli is not None and g.indirimli < g.fiyat for g in gl) / len(gl) if gl else None})
    kategori_ozet = []
    for kat in sorted({u.kategori for u in urunler.values()}, key=katla):
        satir = {"kategori": kat}
        for rk in rakipler:
            e = [s["rakip"][rk]["endeks"] for s in karsilastirma if s["urun"].kategori == kat and rk in s["rakip"]]
            satir[rk] = statistics.median(e) if e else None
        kategori_ozet.append(satir)

    degisim = []
    gruplar = defaultdict(list)
    for g in gozlemler:
        gruplar[(g.rakip, g.kod)].append(g)
    for (rk, kod), lst in gruplar.items():
        lst.sort(key=lambda g: (g.tarih, g.satir))
        if len(lst) >= 2 and lst[0].tarih != lst[-1].tarih:
            ilk, sn = lst[0].fiyat, lst[-1].fiyat
            degisim.append({"rakip": rk, "kod": kod, "urun": urunler[kod].ad, "ilk_tarih": lst[0].tarih, "ilk": ilk, "son_tarih": lst[-1].tarih, "son": sn,
                            "degisim": sn / ilk - 1})
    degisim.sort(key=lambda x: -abs(x["degisim"]))

    bizim_kat = {katla(u.kategori) for u in urunler.values()}
    kamp_ozet = defaultdict(lambda: {"sayi": 0, "indirimler": [], "suren": 0})
    for k in kampanyalar:
        o = kamp_ozet[k["rakip"]]
        o["sayi"] += 1
        if k["indirim"] is not None:
            o["indirimler"].append(k["indirim"])
        k["suren"] = bool(k["bas"] and k["bas"] <= bugun and (k["bit"] is None or k["bit"] >= bugun))
        o["suren"] += k["suren"]
        if k["suren"] and k["indirim"] is not None and k["indirim"] >= derin and (not k["kategori"] or katla(k["kategori"]) in bizim_kat):
            uy.append({"onem": "Orta", "tur": "Süren derin indirim", "kim": k["rakip"],
                       "aciklama": f"{k['kampanya']} — %{k['indirim'] * 100:.0f} ({k['kategori'] or 'tüm ürünler'}), "
                                   f"{k['bas']:%d.%m.%Y}–{k['bit']:%d.%m.%Y}" if k["bit"] else f"{k['kampanya']} — %{k['indirim'] * 100:.0f}, bitiş tarihi yok"})

    # Özellik matrisi
    matris = defaultdict(dict)
    markalar = []
    for o in ozellikler:
        anahtar = (o["kategori"], o["ozellik"])
        matris[anahtar][o["marka"]] = o["deger"]
        if o["marka"] not in markalar:
            markalar.append(o["marka"])
    markalar.sort(key=lambda m: (m != BIZ, katla(m)))
    farklar = []
    for (kat, oz), d in sorted(matris.items(), key=lambda i: (katla(i[0][0]), katla(i[0][1]))):
        biz = d.get(BIZ)
        rakip_deg = [v for m, v in d.items() if m != BIZ]
        var_sayisi = sum(1 for v in rakip_deg if katla(v) in VAR)
        if biz is not None and katla(biz) not in VAR and rakip_deg and var_sayisi >= len(rakip_deg) / 2 and all(katla(v) in VAR | {"yok", "hayir", "h", "0", ""} for v in rakip_deg):
            farklar.append({"kategori": kat, "ozellik": oz, "tur": "Eksik", "aciklama": f"Rakiplerin {var_sayisi}/{len(rakip_deg)}'inde var, bizde yok"})
        elif biz is not None and katla(biz) in VAR and rakip_deg and var_sayisi == 0:
            farklar.append({"kategori": kat, "ozellik": oz, "tur": "Avantaj", "aciklama": "Yalnız bizde var"})
    for f in farklar:
        if f["tur"] == "Eksik":
            uy.append({"onem": "Bilgi", "tur": "Özellik eksiği", "kim": f"{f['kategori']} · {f['ozellik']}", "aciklama": f["aciklama"]})
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    uy.sort(key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))
    return {"urunler": urunler, "gozlemler": gozlemler, "rakipler": rakipler, "karsilastirma": karsilastirma, "rakip_ozet": rakip_ozet, "kategori_ozet": kategori_ozet,
            "degisim": degisim, "kampanyalar": kampanyalar, "kamp_ozet": dict(kamp_ozet), "matris": matris, "markalar": markalar, "farklar": farklar,
            "uyarilar": uy, "bugun": bugun}


def kampanyalari_oku(yol: Path | None) -> list[dict]:
    if not yol:
        return []
    lst = []
    for r in kayitlar(yol, KAMPANYA_SUTUNLARI, ("rakip", "kampanya")):
        if not metin(r.get("rakip")):
            continue
        ind = para(r.get("indirim"))
        lst.append({"bas": tarih(r.get("bas")), "bit": tarih(r.get("bit")), "rakip": metin(r["rakip"]), "kanal": metin(r.get("kanal")), "kampanya": metin(r.get("kampanya")),
                    "indirim": (ind / 100 if ind is not None and ind > 1 else ind), "kategori": metin(r.get("kategori"))})
    return lst


def ozellikleri_oku(yol: Path | None) -> list[dict]:
    if not yol:
        return []
    return [{"marka": metin(r["marka"]), "kategori": metin(r.get("kategori")) or "Genel", "ozellik": metin(r["ozellik"]), "deger": metin(r.get("deger"))}
            for r in kayitlar(yol, OZELLIK_SUTUNLARI, ("marka", "ozellik")) if metin(r.get("marka")) and metin(r.get("ozellik"))]


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
UST = Alignment(vertical="top", wrap_text=True)
PF = "#,##0.00"


def _baslik(ws, basliklar, genislik=()):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "C2"


def _f(x):
    return None if x is None else float(x)


def _endeks_renk(h, e):
    if e is None:
        return
    h.fill = PatternFill("solid", fgColor="E3F4E1" if e >= 105 else "FDE2E1" if e <= 95 else "FFFFFF")


def rapor_yaz(cikti: Path, s: dict) -> None:
    rk = s["rakipler"]
    wb = Workbook()
    oz = wb.active
    oz.title = "Özet"
    _baslik(oz, ["Gösterge", "Değer"], (44, 30))
    oz.freeze_panes = "A2"
    karsi = s["karsilastirma"]
    for a, v in [("Rapor tarihi", s["bugun"].strftime("%d.%m.%Y")), ("Ürünümüz", len(s["urunler"])), ("Rakip", ", ".join(rk)),
                 ("Fiyat gözlemi (eşleşen)", len(s["gozlemler"])), ("En ucuz olduğumuz ürün", sum(x["konum"] == "En ucuz" for x in karsi)),
                 ("En pahalı olduğumuz ürün", sum(x["konum"] == "En pahalı" for x in karsi)), ("Pazar ortası", sum(x["konum"] == "Pazar ortası" for x in karsi)),
                 ("Süren rakip kampanyası", sum(k.get("suren", False) for k in s["kampanyalar"]))]:
        oz.append([a, v])
    oz.append([])
    oz.append(["Endeks = rakibin birim fiyatı / bizim birim fiyatımız × 100. 100'ün üstü: rakip daha pahalı (yeşil ≥ 105), altı: rakip daha ucuz (kırmızı ≤ 95)."])

    fk = wb.create_sheet("Fiyat Karşılaştırma")
    _baslik(fk, ["Ürün Kodu", "Ürün", "Kategori", "Bizim Fiyat", "Bizim Birim Fiyat"] + [f"{r} Endeks" for r in rk] + ["Medyan Endeks", "Konum", "Not"],
            [11, 28, 14, 11, 11] + [11] * len(rk) + [10, 13, 40])
    for x in karsi:
        u = x["urun"]
        fk.append([u.kod, u.ad, u.kategori, float(u.fiyat), float(u.birim_fiyat) if u.miktar else None]
                  + [_f(x["rakip"][r]["endeks"].quantize(Decimal("0.1"))) if r in x["rakip"] else None for r in rk]
                  + [_f(x["medyan"].quantize(Decimal("0.1"))) if x["medyan"] is not None else None, x["konum"], "; ".join(x["notlar"])])
        fk.cell(fk.max_row, 4).number_format = PF
        fk.cell(fk.max_row, 5).number_format = "#,##0.0000"
        for j, r in enumerate(rk, 6):
            _endeks_renk(fk.cell(fk.max_row, j), x["rakip"][r]["endeks"] if r in x["rakip"] else None)
    fk.auto_filter.ref = f"A1:{get_column_letter(fk.max_column)}{fk.max_row}"

    ro = wb.create_sheet("Rakip Özeti")
    _baslik(ro, ["Rakip", "Karşılaştırılan Ürün", "Medyan Endeks", "Bizden Ucuz", "Bizden Pahalı", "Gözlem", "Stokta Olmama", "İndirimli Gözlem", "Kampanya",
                 "Ort. Kampanya İndirimi", "Süren Kampanya"], (18, 10, 10, 9, 9, 8, 10, 10, 9, 11, 9))
    for o in s["rakip_ozet"]:
        k = s["kamp_ozet"].get(o["rakip"], {"sayi": 0, "indirimler": [], "suren": 0})
        ro.append([o["rakip"], o["urun"], _f(o["medyan"].quantize(Decimal("0.1"))) if o["medyan"] is not None else None, o["ucuz"], o["pahali"], o["gozlem"],
                   _f(o["stoksuz"]), _f(o["indirimli"]), k["sayi"], _f(sum(k["indirimler"]) / len(k["indirimler"])) if k["indirimler"] else None, k["suren"]])
        for j in (7, 8, 10):
            ro.cell(ro.max_row, j).number_format = "0%"
        _endeks_renk(ro.cell(ro.max_row, 3), o["medyan"])

    ko = wb.create_sheet("Kategori Özeti")
    _baslik(ko, ["Kategori"] + [f"{r} (medyan endeks)" for r in rk], [18] + [14] * len(rk))
    for x in s["kategori_ozet"]:
        ko.append([x["kategori"]] + [_f(x[r].quantize(Decimal("0.1"))) if x[r] is not None else None for r in rk])
        for j, r in enumerate(rk, 2):
            _endeks_renk(ko.cell(ko.max_row, j), x[r])

    fd = wb.create_sheet("Fiyat Değişimi")
    _baslik(fd, ["Rakip", "Ürün Kodu", "Ürün", "İlk Tarih", "İlk Fiyat", "Son Tarih", "Son Fiyat", "Değişim"], (18, 11, 28, 11, 11, 11, 11, 9))
    for x in s["degisim"]:
        fd.append([x["rakip"], x["kod"], x["urun"], x["ilk_tarih"], float(x["ilk"]), x["son_tarih"], float(x["son"]), float(x["degisim"])])
        fd.cell(fd.max_row, 4).number_format = fd.cell(fd.max_row, 6).number_format = "DD.MM.YYYY"
        fd.cell(fd.max_row, 5).number_format = fd.cell(fd.max_row, 7).number_format = PF
        fd.cell(fd.max_row, 8).number_format = "0.0%"

    ka = wb.create_sheet("Kampanyalar")
    _baslik(ka, ["Başlangıç", "Bitiş", "Rakip", "Kanal", "Kampanya", "Kategori", "İndirim", "Rapor Tarihinde Sürüyor"], (11, 11, 18, 12, 40, 14, 8, 10))
    for k in sorted(s["kampanyalar"], key=lambda k: (k["bas"] or date.min), reverse=True):
        ka.append([k["bas"], k["bit"], k["rakip"], k["kanal"], k["kampanya"], k["kategori"], _f(k["indirim"]), "Evet" if k.get("suren") else ""])
        ka.cell(ka.max_row, 1).number_format = ka.cell(ka.max_row, 2).number_format = "DD.MM.YYYY"
        ka.cell(ka.max_row, 7).number_format = "0%"
        if k.get("suren"):
            ka.cell(ka.max_row, 8).fill = PatternFill("solid", fgColor="FFF4CE")

    om = wb.create_sheet("Özellik Matrisi")
    _baslik(om, ["Kategori", "Özellik"] + s["markalar"] + ["Durum"], [16, 30] + [12] * len(s["markalar"]) + [12])
    fark = {(f["kategori"], f["ozellik"]): f["tur"] for f in s["farklar"]}
    for (kat, oz_), d in sorted(s["matris"].items(), key=lambda i: (katla(i[0][0]), katla(i[0][1]))):
        om.append([kat, oz_] + [d.get(m, "") for m in s["markalar"]] + [fark.get((kat, oz_), "")])
        t = fark.get((kat, oz_))
        if t:
            om.cell(om.max_row, om.max_column).fill = PatternFill("solid", fgColor="FDE2E1" if t == "Eksik" else "E3F4E1")

    ve = wb.create_sheet("Veri")
    _baslik(ve, ["Tarih", "Rakip", "Kanal", "Bizim Ürün Kodu", "Rakip Ürün", "Fiyat", "İndirimli Fiyat", "Ambalaj", "Stokta"], (11, 18, 12, 11, 30, 10, 10, 10, 8))
    for g in sorted(s["gozlemler"], key=lambda g: (g.tarih, g.rakip, g.kod)):
        ve.append([g.tarih, g.rakip, g.kanal, g.kod, g.urun, float(g.fiyat), _f(g.indirimli), f"{g.miktar.normalize():f} {g.temel}" if g.miktar else "",
                   {True: "Evet", False: "Hayır", None: ""}[g.stokta]])
        ve.cell(ve.max_row, 1).number_format = "DD.MM.YYYY"
    ve.auto_filter.ref = f"A1:I{ve.max_row}"

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Ürün / Rakip", "Açıklama"], (9, 24, 34, 100))
    uy.freeze_panes = "A2"
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["kim"], u["aciklama"]])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(urun_yolu: Path, fiyat_yolu: Path, cikti: Path, bugun: date, kampanya_yolu: Path | None = None, ozellik_yolu: Path | None = None,
             esik: Decimal = Decimal("0.10"), derin: Decimal = Decimal("0.30")) -> dict:
    urunler, uy1 = urunleri_oku(urun_yolu)
    if not urunler:
        raise ValueError(f"{urun_yolu.name}: ürün bulunamadı")
    gozlemler, uy2 = gozlemleri_oku(fiyat_yolu, urunler, bugun)
    s = analiz_et(urunler, gozlemler, kampanyalari_oku(kampanya_yolu), ozellikleri_oku(ozellik_yolu), bugun, esik, derin)
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    s["uyarilar"] = sorted(uy1 + uy2 + s["uyarilar"], key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))
    rapor_yaz(cikti, s)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Rakip fiyat, ürün özelliği ve kampanya bilgilerinden karşılaştırmalı rapor hazırlar.")
    p.add_argument("--urunler", type=Path, default=ORNEK / "urunler.csv", help="Ürün Kodu, Ürün, Kategori, Fiyat, Ambalaj Miktarı, Birim")
    p.add_argument("--fiyatlar", type=Path, default=ORNEK / "rakip_fiyatlari.csv",
                   help="Tarih, Rakip, Kanal, Bizim Ürün Kodu, Rakip Ürün, Fiyat, İndirimli Fiyat, Ambalaj Miktarı, Birim, Stokta")
    p.add_argument("--kampanyalar", type=Path, help="Başlangıç, Bitiş, Rakip, Kanal, Kampanya, İndirim (%%), Kategori")
    p.add_argument("--ozellikler", type=Path, help="Marka (bizim için 'Biz'), Kategori, Özellik, Değer (Var / Yok / değer)")
    p.add_argument("--esik", type=float, default=10, help="Fiyat dezavantajı / bırakma eşiği %% (varsayılan 10)")
    p.add_argument("--derin", type=float, default=30, help="Derin indirim eşiği %% (varsayılan 30)")
    p.add_argument("--bugun", help="Rapor tarihi GG.AA.YYYY (örnek veride 09.10.2026)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "rakip_analizi.xlsx")
    a = p.parse_args(argv)
    ornek = a.urunler == ORNEK / "urunler.csv"
    kamp = a.kampanyalar or (ORNEK / "kampanyalar.csv" if ornek else None)
    oz = a.ozellikler or (ORNEK / "ozellikler.csv" if ornek else None)
    for y in (a.urunler, a.fiyatlar, kamp, oz):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    bugun = tarih(a.bugun) if a.bugun else (date(2026, 10, 9) if ornek else date.today())
    if not bugun:
        print("[X] --bugun GG.AA.YYYY biçiminde olmalı")
        return 2
    try:
        s = calistir(a.urunler, a.fiyatlar, a.cikti, bugun, kamp, oz, Decimal(str(a.esik)) / 100, Decimal(str(a.derin)) / 100)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(s['urunler'])} ürün · {len(s['rakipler'])} rakip · {len(s['gozlemler'])} gözlem")
    for o in s["rakip_ozet"]:
        med = "—" if o["medyan"] is None else f"{o['medyan']:.0f}"
        print(f"     {o['rakip']:<20} medyan endeks {med:>4} · bizden ucuz {o['ucuz']}, pahalı {o['pahali']}")
    for u in s["uyarilar"]:
        if u["onem"] != "Bilgi":
            print(f"[!] {u['kim']} · {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
