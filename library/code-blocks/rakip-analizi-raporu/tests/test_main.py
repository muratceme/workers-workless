import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "r.xlsx"
        cls.s = main.calistir(ORNEK / "urunler.csv", ORNEK / "rakip_fiyatlari.csv", cls.cikti, date(2026, 10, 9), ORNEK / "kampanyalar.csv", ORNEK / "ozellikler.csv")
        cls.k = {x["urun"].kod: x for x in cls.s["karsilastirma"]}
        cls.u = {(u["tur"], u["kim"]) for u in cls.s["uyarilar"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_birim_fiyat_endeksi(self):
        x = self.k["D-001"]["rakip"]["Rakip A"]                 # rakip 3 L, biz 2,5 L
        g = x["g"]
        self.assertEqual(g.tarih, date(2026, 10, 5))            # en son gözlem
        beklenen = (g.fiyat / 3000) / (D("289.90") / 2500) * 100
        self.assertEqual(x["endeks"], beklenen)
        b = self.k["D-001"]["rakip"]["Rakip B"]["g"]
        self.assertLess(b.etkin, b.fiyat)                       # indirimli fiyat kullanıldı

    def test_konum_ve_uyarilar(self):
        self.assertEqual(self.k["K-002"]["konum"], "En ucuz")
        self.assertTrue({("Fiyat dezavantajı", "D-002 Toz deterjan 4 kg"), ("Süren derin indirim", "Rakip C (özel marka)"),
                         ("Fiyat bırakma olasılığı", "K-002 Şampuan 500 ml"), ("Eşleşmeyen fiyat", "Rakip A"),
                         ("Özellik eksiği", "Deterjan · Geri dönüştürülmüş ambalaj")} <= self.u)
        self.assertNotIn(("Süren derin indirim", "Rakip B"), self.u)          # %20 < %30
        self.assertEqual({(f["ozellik"], f["tur"]) for f in self.s["farklar"]}, {("Geri dönüştürülmüş ambalaj", "Eksik"), ("Parfümsüz seçenek", "Avantaj")})

    def test_degisim_ve_kampanya(self):
        d = [x for x in self.s["degisim"] if x["rakip"] == "Rakip D"]
        self.assertGreater(sum(x["degisim"] for x in d) / len(d), D("0.02"))   # Rakip D ortalama %4 zam yaptı
        suren = {k["kampanya"] for k in self.s["kampanyalar"] if k["suren"]}
        self.assertEqual(suren, {"Deterjanlarda sezon indirimi", "Toz deterjanda 3 al 2 öde", "Hafta sonu fırsatı (tüm ürünler)"})
        ro = {o["rakip"]: o for o in self.s["rakip_ozet"]}
        self.assertGreater(ro["Rakip D"]["stoksuz"], 0)
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Özet", "Fiyat Karşılaştırma", "Rakip Özeti", "Kategori Özeti", "Fiyat Değişimi", "Kampanyalar", "Özellik Matrisi",
                                         "Veri", "Uyarılar"])


class KuralTesti(unittest.TestCase):
    def test_miktar(self):
        self.assertEqual(main.miktar_coz("1,5", "L"), (D("1500.0"), "ml"))
        self.assertEqual(main.miktar_coz("4 kg", ""), (D(4000), "g"))
        self.assertEqual(main.miktar_coz("750", "ml"), (D(750), "ml"))
        self.assertIsNone(main.miktar_coz("3", "paket"))

    def test_birim_uyusmazligi(self):
        u = main.Urun("A", "x", "K", D(100), D(1000), "ml")
        g = main.Gozlem(1, date(2026, 1, 1), "R", "", "A", "", D(120), None, D(500), "g", None)
        e, n = main.endeks(g, u)
        self.assertEqual(e, D(120))
        self.assertIn("paket fiyatı", n)

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--bugun", "x"]), 2)
            self.assertEqual(main.main(["--fiyatlar", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
