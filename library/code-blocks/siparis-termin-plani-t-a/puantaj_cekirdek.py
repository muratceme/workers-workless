"""
Puantaj Kontrolü — Workers / Workless kod bloğu
İnsan Kaynakları › Bordro ve Özlük İşleri Uzmanı

Aylık puantaj çizelgesini (her gün için bir kod) bordro öncesi kontrol eder:
  - Biçim: tanımsız kod, boş gün, işe giriş öncesi / çıkış sonrası doldurulmuş gün, ayda olmayan gün.
  - Yasal sınırlar (4857 sayılı İş Kanunu): 7 gün üst üste çalışma (md. 46 hafta tatili), günlük 11 saat
    (md. 63 ve Fazla Çalışma Yönetmeliği), yıllık 270 saat fazla çalışma (md. 41).
  - Tatiller: genel tatilde çalışma (md. 47, ek ücret), genel tatil olmayan güne GT kodu, yıllık izne denk gelen
    genel tatil (md. 56: izin süresinden sayılmaz).
  - Devamsızlık olan haftada ödenen hafta tatili (md. 46 koşulu), yıllık izin bakiyesinin aşılması.
  - SGK prim günü ve eksik gün nedeni önerisi; puantajdaki toplam sütunlarıyla (çalışılan gün, fazla mesai,
    SGK gün) karşılaştırma.
İnternete bağlanmaz.

Kullanım:
    python main.py                                         # örnek puantajla dener
    python main.py --girdi puantaj.xlsx --donem 2026-09
    python main.py --girdi puantaj.xlsx --donem 2026-09 --gunluk-saat 7.5 --tatil ek_tatiller.csv
"""
from __future__ import annotations

import argparse
import calendar
import csv
import re
import sys
from collections import Counter
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent

# Puantaj kodları: kanonik kod → (açıklama, ücretli mi, eksik gün nedeni kodu)
KODLAR = {
    "Ç": ("Çalıştı", True, None),
    "HT": ("Hafta tatili", True, None),
    "GT": ("Genel tatil (çalışılmadı)", True, None),
    "Yİ": ("Yıllık ücretli izin", True, None),
    "Mİ": ("Mazeret / yasal ücretli izin (evlilik, ölüm, babalık vb.)", True, None),
    "R": ("Rapor (istirahat)", False, "01"),
    "Üİ": ("Ücretsiz izin", False, "21"),
    "D": ("Devamsızlık", False, "15"),
}
EKSIK_GUN_ADI = {"01": "İstirahat", "15": "Devamsızlık", "21": "Diğer ücretsiz izin", "12": "Birden fazla nedenle"}
ESANLAMLI = {"Ç": ("ç", "c", "x", "1", "çalıştı"), "HT": ("ht", "h", "hafta tatili"), "GT": ("gt", "rt", "bt", "genel tatil"),
             "Yİ": ("yi", "yı", "yıllık izin"), "Mİ": ("mi", "mı", "m", "mazeret"), "R": ("r", "rp", "ri", "rapor"),
             "Üİ": ("üi", "ui", "ü", "ücretsiz izin"), "D": ("d", "ds", "g", "devamsız", "gelmedi")}
ES = {e: k for k, v in ESANLAMLI.items() for e in v}

# Genel tatiller (2912 sayılı Kanun). Dini bayramlar her yıl değişir: Diyanet takvimi; arifeler 13.00'ten itibaren yarım gün.
SABIT_TATILLER = {(1, 1): "Yılbaşı", (4, 23): "Ulusal Egemenlik ve Çocuk Bayramı", (5, 1): "Emek ve Dayanışma Günü",
                  (5, 19): "Atatürk'ü Anma, Gençlik ve Spor Bayramı", (7, 15): "Demokrasi ve Millî Birlik Günü",
                  (8, 30): "Zafer Bayramı", (10, 29): "Cumhuriyet Bayramı"}
DINI_BAYRAMLAR = {  # yıl → [(ilk gün, gün sayısı, ad)]
    2025: [(date(2025, 3, 30), 3, "Ramazan Bayramı"), (date(2025, 6, 6), 4, "Kurban Bayramı")],
    2026: [(date(2026, 3, 20), 3, "Ramazan Bayramı"), (date(2026, 5, 27), 4, "Kurban Bayramı")],
    2027: [(date(2027, 3, 9), 3, "Ramazan Bayramı"), (date(2027, 5, 16), 4, "Kurban Bayramı")],
}


def tatiller(yil: int, ek: dict[date, str] | None = None) -> tuple[dict[date, str], dict[date, str]]:
    """Dönüş: (tam gün genel tatiller, yarım gün arifeler)."""
    tam = {date(yil, a, g): ad for (a, g), ad in SABIT_TATILLER.items()}
    yarim = {date(yil, 10, 28): "Cumhuriyet Bayramı arifesi"}
    for ilk, n, ad in DINI_BAYRAMLAR.get(yil, []):
        yarim[ilk - timedelta(days=1)] = f"{ad} arifesi"
        for i in range(n):
            tam.setdefault(ilk + timedelta(days=i), f"{ad} {i + 1}. gün")
    tam.update(ek or {})
    for d in tam:
        yarim.pop(d, None)
    return tam, yarim


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def katla(s) -> str:
    s = kucuk(s)
    for a, b in zip("ıöüşğç", "iousgc"):
        s = s.replace(a, b)
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def sayi(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    try:
        return Decimal(str(x).strip().replace(".", "").replace(",", ".") if "," in str(x) else str(x).strip())
    except InvalidOperation:
        return None


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for b in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, b).date()
        except ValueError:
            continue
    return None


def hucre_coz(x) -> tuple[str | None, Decimal, str]:
    """'Ç', 'Ç+2', 'X 2,5', 'HT' → (kanonik kod, fazla mesai saati, ham metin). Boşsa kod None, tanımsızsa '?'."""
    ham = str(x if x is not None else "").strip()
    if not ham:
        return None, Decimal(0), ham
    if isinstance(x, (int, float)) and x == 1:
        return "Ç", Decimal(0), ham
    m = re.match(r"^\s*([^\d\s+,.]+|1)\s*\+?\s*(\d+(?:[.,]\d+)?)?\s*$", ham)
    if not m:
        return "?", Decimal(0), ham
    kod = ES.get(kucuk(m.group(1)))
    if not kod:
        return "?", Decimal(0), ham
    fm = Decimal(m.group(2).replace(",", ".")) if m.group(2) else Decimal(0)
    if fm and kod != "Ç":
        return "?", Decimal(0), ham
    return kod, fm, ham


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


AYLAR = ["ocak", "subat", "mart", "nisan", "mayis", "haziran", "temmuz", "agustos", "eylul", "ekim", "kasim", "aralik"]


def donem_bul(satirlar: list[list]) -> tuple[int, int] | None:
    """Başlık üstündeki 'Eylül 2026' / '2026-09' / '09.2026' gibi ifadelerden dönemi bulur."""
    for r in satirlar[:5]:
        metin = katla(" ".join(str(c) for c in r if c))
        m = re.search(r"(20\d\d) (\d{1,2})\b", metin) or None
        if m and 1 <= int(m.group(2)) <= 12:
            return int(m.group(1)), int(m.group(2))
        m = re.search(r"\b(\d{1,2}) (20\d\d)\b", metin)
        if m and 1 <= int(m.group(1)) <= 12:
            return int(m.group(2)), int(m.group(1))
        for i, a in enumerate(AYLAR, 1):
            m = re.search(rf"\b{a}\b\D*(20\d\d)", metin)
            if m:
                return int(m.group(1)), i
    return None


def puantaj_oku(yol: Path) -> tuple[list[dict], tuple[int, int] | None]:
    s = tablo_oku(yol)
    bi = next((i for i, r in enumerate(s[:10]) if sum(1 for c in r if str(c).strip() in ("1", "2", "3", "15", "28")) >= 3), None)
    if bi is None:
        raise SystemExit("Gün sütunları (1, 2, 3 … 31) olan başlık satırı bulunamadı.")
    b = s[bi]
    kb = [katla(x) for x in b]

    def bul(*adlar):
        return next((kb.index(katla(a)) for a in adlar if katla(a) in kb), None)

    gun_sutun = {}
    for i, x in enumerate(b):
        t = str(x).strip()
        if isinstance(x, (int, float)) or t.isdigit():
            g = int(float(t))
            if 1 <= g <= 31 and g not in gun_sutun:
                gun_sutun[g] = i
    i_sicil = bul("sicil", "sicil no", "personel no", "sicil numarası")
    i_ad = bul("ad soyad", "adı soyadı", "personel", "çalışan", "ad")
    i_giris, i_cikis = bul("işe giriş", "giriş tarihi", "işe giriş tarihi"), bul("çıkış", "çıkış tarihi", "işten çıkış")
    beyan = {"calisilan": bul("çalışılan gün", "çalışma günü", "çalışılan"),
             "fm": bul("fazla mesai", "fazla mesai saati", "fm saat", "fm"),
             "sgk": bul("sgk gün", "sgk günü", "prim günü", "sgk prim günü"),
             "kalan_izin": bul("kalan yıllık izin", "kalan izin", "izin bakiyesi"),
             "yil_fm": bul("yıl başından fm", "yıllık fm", "yıl başından fazla mesai", "kümülatif fm")}
    if i_ad is None and i_sicil is None:
        raise SystemExit(f"Personel adı veya sicil sütunu bulunamadı. Başlıklar: {b}")
    personel = []
    for n, r in enumerate(s[bi + 1:], bi + 2):
        al = lambda i: r[i] if i is not None and i < len(r) else None  # noqa: E731
        ad = str(al(i_ad) or "").strip()
        sicil = str(al(i_sicil) or "").strip()
        if not (ad or sicil) or katla(ad).startswith("toplam"):
            continue
        personel.append({"satir": n, "sicil": sicil, "ad": ad or sicil, "giris": tarih(al(i_giris)), "cikis": tarih(al(i_cikis)),
                         "gunler": {g: al(i) for g, i in gun_sutun.items()},
                         "beyan": {k: sayi(al(i)) for k, i in beyan.items() if i is not None}})
    return personel, donem_bul(s[:bi + 1])


# ----------------------------------------------------------------------------
# Kontrol
# ----------------------------------------------------------------------------

def kontrol_et(p: dict, yil: int, ay: int, tam_tatil: dict, yarim_tatil: dict, gunluk_saat: Decimal,
               gunluk_azami: Decimal, yillik_fm: Decimal) -> tuple[dict, list[dict]]:
    ay_gun = calendar.monthrange(yil, ay)[1]
    bas = p["giris"].day if p["giris"] and (p["giris"].year, p["giris"].month) == (yil, ay) else 1
    if p["giris"] and p["giris"] > date(yil, ay, ay_gun):
        bas = ay_gun + 1
    son = p["cikis"].day if p["cikis"] and (p["cikis"].year, p["cikis"].month) == (yil, ay) else ay_gun
    if p["cikis"] and p["cikis"] < date(yil, ay, 1):
        son = 0
    b: list[dict] = []

    def ekle(seviye, kontrol, aciklama, gunler=""):
        b.append({"seviye": seviye, "sicil": p["sicil"], "ad": p["ad"], "kontrol": kontrol, "aciklama": aciklama, "gunler": gunler})

    kodlar: dict[int, str | None] = {}
    fm: dict[int, Decimal] = {}
    tanimsiz, bos, disarida = [], [], []
    for g in range(1, 32):
        kod, saat, ham = hucre_coz(p["gunler"].get(g))
        if g > ay_gun:
            if kod:
                ekle("Hata", "Ayda olmayan gün", f"{g}. gün doldurulmuş ({ham}); bu ay {ay_gun} gün")
            continue
        if bas <= g <= son:
            if kod is None:
                bos.append(g)
            elif kod == "?":
                tanimsiz.append(f"{g}:{ham}")
        elif kod:
            disarida.append(g)
            kod = None
        kodlar[g] = kod if kod != "?" else None
        fm[g] = saat
    if tanimsiz:
        ekle("Hata", "Tanımsız kod", "Tanımsız hücre: " + ", ".join(tanimsiz), ", ".join(t.split(":")[0] for t in tanimsiz))
    if bos:
        ekle("Hata", "Boş gün", f"{len(bos)} gün boş", _gunler(bos))
    if disarida:
        ekle("Hata", "İstihdam dışı gün", "İşe giriş öncesi veya çıkış sonrası günler doldurulmuş (dikkate alınmadı)", _gunler(disarida))

    # 7 gün üst üste çalışma (ay içinde görülebilen)
    seri, ihlal = [], []
    for g in range(1, ay_gun + 1):
        if kodlar.get(g) == "Ç":
            seri.append(g)
            if len(seri) == 7:
                ihlal.append(seri[0])
        else:
            seri = []
    if ihlal:
        ekle("Yasal", "Hafta tatili", "7 gün üst üste çalışma: 7 günlük dönemde en az 24 saat kesintisiz dinlenme verilmeli "
             "(İş K. md. 46). Hafta tatilinde çalıştırma varsa ayrıca ücretlendirilmeli.",
             ", ".join(f"{g}-{g + 6}" for g in ihlal))
    # Günlük 11 saat
    asan = [g for g, s in fm.items() if s and gunluk_saat + s > gunluk_azami]
    if asan:
        ekle("Yasal", "Günlük 11 saat", f"Normal {_s(gunluk_saat)} saat + fazla mesai, günlük {_s(gunluk_azami)} saat sınırını aşıyor "
             "(İş K. md. 63, Fazla Çalışma Yönetmeliği)", _gunler(asan))
    fm_toplam = sum(fm.values(), Decimal(0))
    yil_fm = p["beyan"].get("yil_fm")
    if yil_fm is not None and yil_fm + fm_toplam > yillik_fm:
        ekle("Yasal", "Yıllık 270 saat", f"Yıl başından {_s(yil_fm)} + bu ay {_s(fm_toplam)} = {_s(yil_fm + fm_toplam)} saat; "
             f"yıllık sınır {_s(yillik_fm)} saat (İş K. md. 41)")

    # Tatiller
    gt_calisma = [g for g in kodlar if kodlar[g] == "Ç" and date(yil, ay, g) in tam_tatil]
    if gt_calisma:
        ekle("Bilgi", "Genel tatilde çalışma", "Genel tatilde çalışılan her gün için ayrıca bir günlük ücret ödenir (İş K. md. 47): "
             + ", ".join(f"{g} ({tam_tatil[date(yil, ay, g)]})" for g in gt_calisma), _gunler(gt_calisma))
    yanlis_gt = [g for g in kodlar if kodlar[g] == "GT" and date(yil, ay, g) not in tam_tatil]
    if yanlis_gt:
        ekle("Hata", "GT kodu", "Genel tatil olmayan güne GT kodu verilmiş"
             + (" (arife yarım gündür; sabah çalışılır)" if any(date(yil, ay, g) in yarim_tatil for g in yanlis_gt) else ""),
             _gunler(yanlis_gt))
    izin_gt = [g for g in kodlar if kodlar[g] == "Yİ" and date(yil, ay, g) in tam_tatil]
    if izin_gt:
        ekle("Dikkat", "İzinde genel tatil", "Yıllık izne denk gelen genel tatil izin süresinden sayılmaz (İş K. md. 56); "
             "GT olarak kodlayın", _gunler(izin_gt))

    # Devamsızlık olan haftada hafta tatili
    for pzt in sorted({date(yil, ay, g) - timedelta(days=date(yil, ay, g).weekday()) for g in kodlar}):
        hafta = [(pzt + timedelta(days=i)) for i in range(7)]
        gunler = [d.day for d in hafta if d.month == ay and d.day in kodlar]
        d_gun = [g for g in gunler if kodlar[g] == "D"]
        ht_gun = [g for g in gunler if kodlar[g] == "HT"]
        if d_gun and ht_gun:
            ekle("Dikkat", "Devamsızlık ve hafta tatili",
                 f"{_gunler(d_gun)} devamsızlık olan haftada {_gunler(ht_gun)} hafta tatili ücretli kodlanmış. Hafta tatili ücreti, "
                 "tatilden önceki iş günlerinde çalışmış olma koşuluna bağlıdır (İş K. md. 46); işyeri uygulamanızı kontrol edin.",
                 _gunler(d_gun + ht_gun))

    say = Counter(k for k in kodlar.values() if k)
    kalan = p["beyan"].get("kalan_izin")
    if kalan is not None and say["Yİ"] > kalan:
        ekle("Dikkat", "İzin bakiyesi", f"Bu ay {say['Yİ']} gün yıllık izin kullanılmış; kalan bakiye {_s(kalan)} gün")

    # SGK prim günü
    istihdam = max(0, son - bas + 1)
    tam_ay = bas == 1 and son == ay_gun
    taban = 30 if tam_ay else min(30, istihdam)
    eksik = {neden: sum(1 for k in kodlar.values() if k and KODLAR[k][2] == neden) for neden in ("01", "15", "21")}
    eksik_toplam = sum(eksik.values())
    sgk = max(0, taban - eksik_toplam)
    nedenler = [n for n, v in eksik.items() if v]
    neden_kodu = (nedenler[0] if len(nedenler) == 1 else "12") if nedenler else ""

    ozet = {"sicil": p["sicil"], "ad": p["ad"], "kodlar": kodlar, "fm_gun": fm, "sayac": say, "fm": fm_toplam,
            "istihdam": istihdam, "sgk": sgk, "eksik": eksik_toplam, "eksik_kodu": neden_kodu, "ay_gun": ay_gun}
    # Beyan edilen toplamlarla karşılaştırma
    bey = p["beyan"]
    if bey.get("calisilan") is not None and bey["calisilan"] != say["Ç"]:
        ekle("Hata", "Toplam: çalışılan gün", f"Puantajda {say['Ç']} gün Ç, toplam sütununda {_s(bey['calisilan'])}")
    if bey.get("fm") is not None and bey["fm"] != fm_toplam:
        ekle("Hata", "Toplam: fazla mesai", f"Günlük fazla mesai toplamı {_s(fm_toplam)} saat, toplam sütununda {_s(bey['fm'])}")
    if bey.get("sgk") is not None and bey["sgk"] != sgk:
        fark = abs(bey["sgk"] - sgk)
        ay_notu = " (31/28/29 çeken ayda gün hesabı uygulaması farkı olabilir)" if fark == 1 and ay_gun != 30 and bey["sgk"] <= 30 else ""
        if bey["sgk"] > 30:
            ay_notu = ""
            ekle("Hata", "SGK gün > 30", "Aylık prim günü 30'dan fazla bildirilemez")
        ekle("Dikkat" if ay_notu else "Hata", "Toplam: SGK gün",
             f"Hesaplanan prim günü {sgk} (taban {taban} − eksik {eksik_toplam}), toplam sütununda {_s(bey['sgk'])}{ay_notu}")
    return ozet, b


def _s(x) -> str:
    """Decimal → Türkçe sayı metni: 7.5 → '7,5', 270.0 → '270'."""
    return f"{Decimal(x).normalize():f}".replace(".", ",")


def _gunler(gs: list[int]) -> str:
    """[1,2,3,7] → '1-3, 7'"""
    gs = sorted(set(gs))
    parca, i = [], 0
    while i < len(gs):
        j = i
        while j + 1 < len(gs) and gs[j + 1] == gs[j] + 1:
            j += 1
        parca.append(f"{gs[i]}" if i == j else f"{gs[i]}-{gs[j]}")
        i = j + 1
    return ", ".join(parca)


def ek_tatil_oku(yol: Path | None) -> dict[date, str]:
    if not yol:
        return {}
    sonuc = {}
    for r in tablo_oku(yol):
        d = tarih(r[0])
        if d:
            sonuc[d] = str(r[1]).strip() if len(r) > 1 and r[1] else "Ek tatil"
    return sonuc


SEVIYE_SIRA = {"Hata": 0, "Yasal": 1, "Dikkat": 2, "Bilgi": 3}


def calistir(girdi: Path, cikti: Path, donem: tuple[int, int] | None = None, gunluk_saat: float = 7.5,
             gunluk_azami: float = 11.0, yillik_fm: float = 270.0, ek_tatil: Path | None = None) -> dict:
    personel, bulunan = puantaj_oku(girdi)
    donem = donem or bulunan
    if not donem:
        raise SystemExit("Dönem bulunamadı: --donem YYYY-AA verin (ör. --donem 2026-09).")
    yil, ay = donem
    uyarilar = []
    if yil not in DINI_BAYRAMLAR:
        uyarilar.append(f"{yil} yılı dini bayram tarihleri tanımlı değil; --tatil ile verin.")
    tam, yarim = tatiller(yil, ek_tatil_oku(ek_tatil))
    ozetler, bulgular = [], []
    for p in personel:
        o, b = kontrol_et(p, yil, ay, tam, yarim, Decimal(str(gunluk_saat)), Decimal(str(gunluk_azami)), Decimal(str(yillik_fm)))
        ozetler.append(o)
        bulgular += b
    bulgular.sort(key=lambda x: (SEVIYE_SIRA[x["seviye"]], x["sicil"], x["kontrol"]))
    ay_tatil = {d.day: ad for d, ad in tam.items() if (d.year, d.month) == (yil, ay)}
    ay_yarim = {d.day: ad for d, ad in yarim.items() if (d.year, d.month) == (yil, ay)}
    _rapor(ozetler, bulgular, yil, ay, ay_tatil, ay_yarim, uyarilar, cikti)
    return {"ozetler": ozetler, "bulgular": bulgular, "donem": donem, "tatiller": ay_tatil, "uyarilar": uyarilar}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Hata": "F8C9C6", "Yasal": "FDE2E1", "Dikkat": "FFF4CE", "Bilgi": "E8F0FE"}
KOD_RENK = {"HT": "E5E5EA", "GT": "D1E7FF", "Yİ": "E3F5E1", "Mİ": "E3F5E1", "R": "FFE8CC", "Üİ": "FFF4CE", "D": "F8C9C6"}
GUN_ADI = ["Pt", "Sa", "Ça", "Pe", "Cu", "Ct", "Pz"]


def _baslik(ws, satir=1):
    for c in ws[satir]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI


def _rapor(ozetler, bulgular, yil, ay, ay_tatil, ay_yarim, uyarilar, cikti):
    wb = Workbook()
    o = wb.active
    o.title = "Özet"
    o.append([f"Puantaj kontrolü · {AYLAR[ay - 1].capitalize().replace('Subat', 'Şubat').replace('Mayis', 'Mayıs').replace('Agustos', 'Ağustos').replace('Eylul', 'Eylül').replace('Kasim', 'Kasım').replace('Aralik', 'Aralık')} {yil}"])
    o["A1"].font = Font(bold=True, size=12)
    o.append(["Personel", len(ozetler)])
    o.append(["Bulgulu personel", len({(b["sicil"], b["ad"]) for b in bulgular if b["seviye"] != "Bilgi"})])
    for s in SEVIYE_SIRA:
        o.append([s, sum(1 for b in bulgular if b["seviye"] == s)])
        o.cell(o.max_row, 1).fill = PatternFill("solid", fgColor=RENK[s])
    o.append([])
    o.append(["Bu aydaki genel tatiller"])
    o.cell(o.max_row, 1).font = Font(bold=True)
    for g, ad in sorted(ay_tatil.items()):
        o.append([f"{g:02d}.{ay:02d}.{yil}", ad])
    for g, ad in sorted(ay_yarim.items()):
        o.append([f"{g:02d}.{ay:02d}.{yil}", ad + " (13.00'ten itibaren yarım gün)"])
    if not ay_tatil and not ay_yarim:
        o.append(["-", "Yok"])
    for u in uyarilar:
        o.append(["Uyarı", u])
    o.append([])
    o.append(["Hata: puantaj veya toplamlar hatalı · Yasal: İş Kanunu sınırı aşılmış · Dikkat: bordroda karar gerektirir · Bilgi: ek ödeme vb."])
    o.column_dimensions["A"].width = 22
    o.column_dimensions["B"].width = 60

    b = wb.create_sheet("Bulgular")
    b.append(["Seviye", "Sicil", "Ad Soyad", "Kontrol", "Günler", "Açıklama", "Düzeltildi / Not"])
    _baslik(b)
    for x in bulgular:
        b.append([x["seviye"], x["sicil"], x["ad"], x["kontrol"], x["gunler"], x["aciklama"], None])
        b.cell(b.max_row, 1).fill = PatternFill("solid", fgColor=RENK[x["seviye"]])
        b.cell(b.max_row, 6).alignment = Alignment(wrap_text=True, vertical="top")
    for j, w in enumerate((8, 10, 22, 24, 14, 100, 24), 1):
        b.column_dimensions[get_column_letter(j)].width = w
    b.freeze_panes = "D2"
    b.auto_filter.ref = b.dimensions

    p = wb.create_sheet("Personel Özeti")
    kodlar = list(KODLAR)
    p.append(["Sicil", "Ad Soyad", "İstihdam Günü"] + kodlar + ["Fazla Mesai (saat)", "Eksik Gün", "Eksik Gün Nedeni", "SGK Prim Günü"])
    _baslik(p)
    for z in ozetler:
        p.append([z["sicil"], z["ad"], z["istihdam"]] + [z["sayac"].get(k, 0) for k in kodlar] +
                 [float(z["fm"]), z["eksik"], f"{z['eksik_kodu']} {EKSIK_GUN_ADI.get(z['eksik_kodu'], '')}".strip(), z["sgk"]])
    for j, w in enumerate([10, 22, 10] + [6] * len(kodlar) + [12, 9, 22, 10], 1):
        p.column_dimensions[get_column_letter(j)].width = w
    p.freeze_panes = "C2"

    c = wb.create_sheet("Puantaj")
    ay_gun = calendar.monthrange(yil, ay)[1]
    c.append(["Sicil", "Ad Soyad"] + list(range(1, ay_gun + 1)))
    c.append(["", ""] + [GUN_ADI[date(yil, ay, g).weekday()] for g in range(1, ay_gun + 1)])
    _baslik(c)
    for g in range(1, ay_gun + 1):
        if g in ay_tatil:
            c.cell(2, g + 2).fill = PatternFill("solid", fgColor="3A6FB0")
    for z in ozetler:
        c.append([z["sicil"], z["ad"]] + [(z["kodlar"].get(g) or "") + (f"+{z['fm_gun'][g]:g}" if z["fm_gun"].get(g) else "")
                                          for g in range(1, ay_gun + 1)])
        for g in range(1, ay_gun + 1):
            k = z["kodlar"].get(g)
            if k in KOD_RENK:
                c.cell(c.max_row, g + 2).fill = PatternFill("solid", fgColor=KOD_RENK[k])
    c.column_dimensions["B"].width = 22
    for g in range(1, ay_gun + 1):
        c.column_dimensions[get_column_letter(g + 2)].width = 5
    c.freeze_panes = "C3"

    k = wb.create_sheet("Kodlar")
    k.append(["Kod", "Kabul edilen yazımlar", "Açıklama", "Ücretli", "SGK eksik gün nedeni"])
    _baslik(k)
    for kod, (ac, ucretli, neden) in KODLAR.items():
        k.append([kod, ", ".join(ESANLAMLI[kod]), ac, "Evet" if ucretli else "Hayır",
                  f"{neden} {EKSIK_GUN_ADI[neden]}" if neden else ""])
    k.append([])
    for s in ["Fazla mesai: çalışılan güne saat eklenir, ör. Ç+2 veya X 2,5",
              "SGK prim günü: tam ayda 30 − eksik gün; ay içinde giriş/çıkışta istihdam edilen gün (en çok 30) − eksik gün",
              "Birden fazla eksik gün nedeni varsa SGK'da 12 (Birden fazla) kodu kullanılır; belgeler (rapor, tutanak, izin "
              "dilekçesi) saklanmalıdır",
              "Rapor: geçici iş göremezlik ödeneği istirahatin 3. gününden itibaren SGK tarafından ödenir (5510 md. 18)",
              "Bulgular kontrol önerisidir; bordro ve SGK bildirimi öncesi uzman kontrolü yapın"]:
        k.append(["Not", s])
    k.column_dimensions["A"].width = 6
    k.column_dimensions["B"].width = 34
    k.column_dimensions["C"].width = 60
    k.column_dimensions["E"].width = 24
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description="Aylık puantajı bordro öncesi tutarlılık ve İş Kanunu sınırları açısından kontrol eder.")
    ap.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "puantaj_mayis_2026.csv",
                    help="Puantaj (.xlsx/.csv): Sicil, Ad Soyad, [İşe Giriş, Çıkış], 1 … 31, [Çalışılan Gün, Fazla Mesai, SGK Gün, ...]")
    ap.add_argument("--donem", help="YYYY-AA (dosya başlığında 'Eylül 2026' gibi yazıyorsa gerekmez)")
    ap.add_argument("--gunluk-saat", type=float, default=7.5, help="Günlük normal çalışma saati (varsayılan 7,5 = 45/6)")
    ap.add_argument("--gunluk-azami", type=float, default=11.0, help="Günlük azami çalışma (varsayılan 11)")
    ap.add_argument("--yillik-fm", type=float, default=270.0, help="Yıllık fazla çalışma sınırı (varsayılan 270)")
    ap.add_argument("--tatil", type=Path, help="Ek tatil günleri (.csv/.xlsx: Tarih, Açıklama) — idari izin, yeni yıl bayram tarihleri")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "puantaj_kontrolu.xlsx")
    a = ap.parse_args(argv)
    donem = None
    if a.donem:
        m = re.match(r"^(\d{4})-(\d{1,2})$", a.donem)
        if not m:
            raise SystemExit("--donem biçimi YYYY-AA olmalı (ör. 2026-09)")
        donem = (int(m.group(1)), int(m.group(2)))
    s = calistir(a.girdi, a.cikti, donem, a.gunluk_saat, a.gunluk_azami, a.yillik_fm, a.tatil)
    sayac = Counter(b["seviye"] for b in s["bulgular"])
    print(f"[OK] {s['donem'][0]}-{s['donem'][1]:02d} · {len(s['ozetler'])} personel · "
          + " · ".join(f"{k}: {sayac.get(k, 0)}" for k in SEVIYE_SIRA))
    for u in s["uyarilar"]:
        print(f"[!] {u}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
