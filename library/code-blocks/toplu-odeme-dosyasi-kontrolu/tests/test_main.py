import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
# Belgelerde kullanılan örnek IBAN'lar (gerçek hesaba ait değildir)
TR1 = "TR330006100519786457841326"
GB1 = "GB82WEST12345698765432"
DE1 = "DE89370400440532013000"


def iban_uret(banka: str, hesap: int) -> str:
    govde = f"{banka}0{hesap:016d}"
    return f"TR{98 - int(govde + '292700') % 97:02d}{govde}"


TR2, TR3 = iban_uret("00062", 123456789), iban_uret("00010", 987654321)


def satir(ad="Ali Veli", iban=TR1, tutar="1.000,00", kimlik="", n=2):
    return {"satir": n, "ad": ad, "iban_ham": iban, "tutar_ham": tutar, "kimlik": "".join(c for c in kimlik if c.isdigit()), "kimlik_ham": kimlik, "aciklama": ""}


def turler(s):
    return sorted(t for _, t, _ in s["bulgular"])


class IbanTesti(unittest.TestCase):
    def test_mod97(self):
        for i in (TR1, GB1, DE1, TR2, TR3):
            self.assertTrue(main.iban_mod97(i), i)
        self.assertFalse(main.iban_mod97(TR1[:-1] + "7"))
        self.assertFalse(main.iban_mod97("TR33"))

    def test_tr_kurallari(self):
        self.assertEqual(main.iban_kontrol("tr33 0006 1005 1978 6457 8413 26"), (TR1, None))
        self.assertIn("26 karakter", main.iban_kontrol(TR1[:-1])[1])
        self.assertIn("mod 97", main.iban_kontrol(TR1[:-2] + "62")[1])
        self.assertIn("yalnız rakam", main.iban_kontrol(TR1[:12] + "A" + TR1[13:])[1])
        self.assertIn("rezerv", main.iban_kontrol(TR1[:10] + "1" + TR1[11:])[1])
        self.assertEqual(main.iban_kontrol("")[1], "IBAN boş")
        self.assertIn("ülke koduyla", main.iban_kontrol("1234567890")[1])
        self.assertIsNone(main.iban_kontrol(DE1)[1])

    def test_kimlik_no(self):
        self.assertTrue(main.tckn_gecerli("10000000146"))                           # bilinen test numarası
        self.assertFalse(main.tckn_gecerli("10000000147"))
        self.assertFalse(main.tckn_gecerli("00000000000"))
        self.assertTrue(main.vkn_gecerli("1234567890"))                             # algoritmaya göre geçerli örnek
        self.assertFalse(main.vkn_gecerli("1234567891"))


class KontrolTesti(unittest.TestCase):
    def test_temiz(self):
        r = main.kontrol([satir(), satir("Ayşe Kaya", TR2, "2500,50", n=3)], beklenen_adet=2, beklenen_toplam=D("3500.50"))
        self.assertEqual((r["sonuc"], r["toplam"], r["genel"]), ("Biçim kontrolünden geçti", D("3500.50"), []))
        self.assertEqual(r["banka"], {"Garanti BBVA": [1, D("2500.50")], "Diğer / tanımsız": [1, D("1000.00")]})

    def test_satir_hatalari(self):
        s = [satir(tutar="0"), satir(tutar="12,345", n=3), satir(tutar="bin", n=4), satir(ad="", iban=TR2, n=5), satir(iban=DE1, n=6), satir(iban=TR3, kimlik="10000000147", n=7),
             satir(iban=iban_uret("00064", 5), kimlik="123", n=8)]
        r = main.kontrol(s, gonderen=TR2, ust_sinir=D(500))
        self.assertEqual([turler(x) for x in r["satirlar"]], [["Tutar sıfır veya eksi"], ["Tutar 2 ondalıktan fazla"], ["Tutar okunamadı"],
                                                             ["Alıcı adı eksik", "Gönderen hesabına ödeme", "Tek ödeme sınırı aşıldı"], ["Tek ödeme sınırı aşıldı", "Yabancı IBAN"],
                                                             ["TCKN hatalı", "Tek ödeme sınırı aşıldı"], ["Kimlik no uzunluğu hatalı", "Tek ödeme sınırı aşıldı"]])
        self.assertEqual([x["durum"] for x in r["satirlar"]], ["Hatalı", "Hatalı", "Hatalı", "Hatalı", "Kontrol edin", "Kontrol edin", "Kontrol edin"])
        self.assertEqual((r["sonuc"], r["temiz_adet"]), ("İşleme alınamaz", 3))

    def test_mukerrer(self):
        s = [satir(), satir(n=3), satir("Başka Biri", TR1, "750,00", n=4), satir("Ayşe Kaya", TR2, "100", "10000000146", 5), satir("Ayşe Kaya", TR3, "200", "10000000146", 6)]
        r = main.kontrol(s)
        self.assertEqual([turler(x) for x in r["satirlar"]], [["Aynı IBAN, farklı alıcı"], ["Aynı IBAN, farklı alıcı", "Mükerrer ödeme"], ["Aynı IBAN, farklı alıcı"],
                                                             ["Aynı kimlik no, farklı IBAN"], ["Aynı kimlik no, farklı IBAN"]])
        self.assertIn("satır 2", r["satirlar"][1]["bulgular"][0][2])

    def test_kontrol_toplami(self):
        r = main.kontrol([satir(), satir(iban=TR2, n=3)], dip=D("2000.00"), beklenen_adet=3, beklenen_toplam=D("2500"))
        self.assertEqual([(u["kim"], u["aciklama"]) for u in r["genel"]],
                         [("Adet", "Beklenen 3 · dosyada 2 · fark -1"), ("Toplam tutar", "Beklenen 2.500,00 · dosyada 2.000,00 · fark -500,00")])
        self.assertEqual(len(main.kontrol([satir()], dip=D("999"))["genel"]), 1)

    def test_onceki_dosya(self):
        onceki = [satir("Ali Veli", TR1, "1000", "10000000146"), satir("Ayşe Kaya", TR2, "2000"), satir("Giden Kişi", TR3, "500")]
        s = [satir("ALİ VELİ", TR3, "1.300,00", "10000000146"), satir("Ayşe Kaya", TR2, "2.400,00", n=3), satir("Yeni Kişi", iban_uret("00046", 77), "900", n=4),
             satir("Ayşe Kaya", TR2[:-1], "2000", n=5)]
        r = main.kontrol(s, onceki=onceki)
        self.assertEqual([turler(x) for x in r["satirlar"]], [["IBAN değişmiş", "Tutar belirgin değişmiş"], [], ["Yeni alıcı"], ["IBAN hatalı"]])   # %30 > 25; %20 uyarı vermez
        self.assertEqual([u["aciklama"] for u in r["genel"]], ["Giden Kişi"])
        ayni = main.kontrol([satir("Ali Veli", TR1, "1000"), satir("Ayşe Kaya", TR2, "2000", n=3)], onceki=[satir("Ali Veli", TR1, "1000"), satir("Ayşe Kaya", TR2, "2000")])
        self.assertEqual([u["tur"] for u in ayni["genel"]], ["Dosya önceki dosyanın birebir aynısı"])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "odemeler.csv", cikti, ORNEK / "onceki_ay.csv", beklenen_adet=40, beklenen_toplam=D("1702740.00"))
            self.assertEqual((r["adet"], r["toplam"], r["dip"]), (41, D("1750540.00"), D("1750540.00")))
            hatali = {s["satir"]: turler(s) for s in r["satirlar"] if s["durum"] == "Hatalı"}
            self.assertEqual(hatali, {6: ["IBAN hatalı"], 11: ["IBAN hatalı"], 14: ["Mükerrer ödeme"], 17: ["IBAN değişmiş"], 26: ["Tutar sıfır veya eksi"]})
            self.assertEqual(sum(t == "Yeni alıcı" for s in r["satirlar"] for t in turler(s)), 2)
            self.assertEqual([u["tur"] for u in r["genel"]], ["Kontrol toplamı tutmuyor", "Kontrol toplamı tutmuyor", "Önceki dosyada olup bu dosyada olmayan alıcı"])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Satır Kontrolü", "Bulgular"])

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp) / "o.csv"
            d.write_text(f"Alıcı;IBAN;Tutar\nAli Veli;{TR1};1.000,00\nAyşe Kaya;{TR2};500,00\nTOPLAM;;1.500,00\n", encoding="utf-8")
            arg = ["--dosya", str(d), "--cikti", str(Path(tmp) / "x.xlsx")]
            self.assertEqual(main.main(arg + ["--beklenen-adet", "2", "--beklenen-toplam", "1.500,00"]), 0)
            self.assertEqual(main.oku(d)[1], D("1500.00"))
            self.assertEqual(main.main(arg + ["--beklenen-toplam", "çok"]), 2)
            self.assertEqual(main.main(["--dosya", str(Path(tmp) / "yok.csv")]), 1)
            d.write_text("Alıcı;IBAN;Tutar\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "y.xlsx")]), 0)


if __name__ == "__main__":
    unittest.main()
