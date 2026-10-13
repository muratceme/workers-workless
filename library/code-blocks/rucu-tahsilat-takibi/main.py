"""
Rücu Tahsilat Takibi — Workers / Workless kod bloğu
Hasar Yönetimi › Rücu Uzmanı

Rücu dosyalarını talep, cevap, tahkim/dava ve tahsilat aşamalarında izler; bekleyen ve süresi daralan dosyaları öne çıkarır.
  Kalan tutar          = rücu tutarı − tahsil edilen;   tahsilat oranı % = tahsil edilen ÷ rücu tutarı × 100
  Kapalı dosya         = aşaması "Tahsil edildi" ya da "Kapatıldı" olan veya tamamı tahsil edilen dosya
  Zamanaşımı geçmiş    = zamanaşımı tarihi < durum tarihi;   yaklaşıyor = tarihe --gun (90) gün veya az kaldı
  Talep gönderilmedi   = talep tarihi boş
  Cevap gecikti        = talep gönderilmiş, cevap yok ve --cevap (30) günden fazla geçmiş
  Ret, karar bekliyor  = aşama "Reddedildi" (tahkim ya da dava kararı verilecek)
  Tahsilat gecikti     = aşama "Kabul edildi", tutar kalmış ve cevabın üzerinden --tahsil (30) günden fazla geçmiş
Zamanaşımı tarihini siz yazarsınız; kod süre hesaplamaz, hukuki değerlendirme yapmaz.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"no": ("rucu no", "rucu dosya no", "no"), "hasar": ("hasar dosya no", "hasar no", "dosya no"), "karsi": ("karsi taraf", "muhatap", "borclu"), "tutar": ("rucu tutari", "talep tutari", "tutar"),
         "talep": ("talep tarihi", "gonderim tarihi", "rucu talep tarihi"), "cevap": ("cevap tarihi", "yanit tarihi"), "asama": ("asama", "durum"), "tahsil": ("tahsil edilen", "tahsilat", "tahsil edilen tutar"),
         "zaman": ("zamanasimi tarihi", "zamanasimi", "son tarih"), "sorumlu": ("sorumlu", "rucu uzmani")}
ASAMA = {"hazirlik": "Hazırlık", "talep gonderildi": "Talep gönderildi", "gonderildi": "Talep gönderildi", "cevap bekleniyor": "Talep gönderildi", "kabul edildi": "Kabul edildi", "kabul": "Kabul edildi",
         "reddedildi": "Reddedildi", "ret": "Reddedildi", "tahkim": "Tahkim", "dava": "Dava / İcra", "icra": "Dava / İcra", "dava icra": "Dava / İcra", "tahsil edildi": "Tahsil edildi", "tahsil": "Tahsil edildi",
         "kapatildi": "Kapatıldı", "kapandi": "Kapatıldı", "vazgecildi": "Kapatıldı"}
ASAMA_SIRA = ("Hazırlık", "Talep gönderildi", "Kabul edildi", "Reddedildi", "Tahkim", "Dava / İcra", "Tahsil edildi", "Kapatıldı")
DURUM = ("Zamanaşımı geçmiş", "Zamanaşımı yaklaşıyor", "Talep gönderilmedi", "Cevap gecikti", "Ret, karar bekliyor", "Tahsilat gecikti", "Süreçte")
DURUM_RENK = {"Zamanaşımı geçmiş": KIRMIZI, "Zamanaşımı yaklaşıyor": KIRMIZI, "Talep gönderilmedi": SARI, "Cevap gecikti": SARI, "Ret, karar bekliyor": SARI, "Tahsilat gecikti": SARI, "Süreçte": None,
              "Tahsil edildi": YESIL, "Kısmen tahsil edildi": MAVI, "Kapatıldı": MAVI}


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("no", "tutar")):
        no = metin(r["no"])
        if not no:
            continue
        tutar, tahsil = para(r["tutar"]), para(r["tahsil"]) if metin(r["tahsil"]) else Decimal(0)
        if tutar is None or tutar <= 0 or tahsil is None or tahsil < 0 or katla(no) in gorulen:
            uy.append(uyari("Yüksek", "Dosya okunamadı", no, f"Satır {r['_satir']}: " + ("aynı rücu numarası daha önce geçti" if katla(no) in gorulen else f"rücu tutarı '{metin(r['tutar'])}' ya da tahsil edilen '{metin(r['tahsil'])}' okunamadı")
                            + ". Satır alınmadı."))
            continue
        gorulen.add(katla(no))
        for alan, ad in (("talep", "Talep tarihi"), ("cevap", "Cevap tarihi"), ("zaman", "Zamanaşımı tarihi")):
            if metin(r[alan]) and tarih(r[alan]) is None:
                uy.append(uyari("Orta", "Tarih okunamadı", no, f"{ad} '{metin(r[alan])}' okunamadı; boş sayıldı."))
        asama = ASAMA.get(katla(r["asama"]))
        if asama is None and metin(r["asama"]):
            uy.append(uyari("Bilgi", "Aşama tanınmadı", no, f"'{metin(r['asama'])}' bilinen aşamalardan biri değil ({', '.join(ASAMA_SIRA)}); olduğu gibi kullanıldı."))
        l.append({"no": no, "hasar": metin(r["hasar"]), "karsi": metin(r["karsi"]) or "(karşı taraf yok)", "tutar": tutar, "talep": tarih(r["talep"]), "cevap": tarih(r["cevap"]), "asama": asama or metin(r["asama"]) or "Hazırlık",
                  "tahsil": tahsil, "zaman": tarih(r["zaman"]), "sorumlu": metin(r["sorumlu"])})
    return l, uy


def _oran(pay: Decimal, payda: Decimal) -> Decimal | None:
    return yuvarla(pay * 100 / payda) if payda else None


def degerlendir(dosyalar: list[dict], bugun: date, gun: int = 90, cevap: int = 30, tahsil: int = 30) -> dict:
    if not dosyalar:
        raise ValueError("İzlenecek rücu dosyası yok.")
    uy = []
    for d in dosyalar:
        kim = f"{d['no']} · {d['karsi']}"
        d["kalan"] = max(d["tutar"] - d["tahsil"], Decimal(0))
        d["oran"] = _oran(d["tahsil"], d["tutar"])
        d["cevap_sure"] = (d["cevap"] - d["talep"]).days if d["talep"] and d["cevap"] else None
        d["bekleme"] = (bugun - d["talep"]).days if d["talep"] and not d["cevap"] else None
        d["zaman_kalan"] = (d["zaman"] - bugun).days if d["zaman"] else None
        d["acik"] = not (d["asama"] in ("Tahsil edildi", "Kapatıldı") or d["kalan"] == 0)
        if d["tahsil"] > d["tutar"]:
            uy.append(uyari("Orta", "Tahsilat rücu tutarından büyük", kim, f"Rücu tutarı {d['tutar']:,.2f}, tahsil edilen {d['tahsil']:,.2f}. Tutarları kontrol edin (faiz ve masraf ayrı izlenmeli)."))
        if not d["acik"]:
            d["durum"] = "Tahsil edildi" if d["kalan"] == 0 else "Kısmen tahsil edildi" if d["asama"] == "Tahsil edildi" else "Kapatıldı"
            continue
        bilgi = f"Kalan {d['kalan']:,.2f} · aşama {d['asama']}."
        if d["zaman_kalan"] is not None and d["zaman_kalan"] < 0:
            d["durum"] = "Zamanaşımı geçmiş"
            uy.append(uyari("Yüksek", "Zamanaşımı tarihi geçmiş", kim, f"Zamanaşımı tarihi {d['zaman']:%d.%m.%Y} ({-d['zaman_kalan']} gün önce). {bilgi} Dosyayı hukukla değerlendirin."))
        elif d["zaman_kalan"] is not None and d["zaman_kalan"] <= gun:
            d["durum"] = "Zamanaşımı yaklaşıyor"
            uy.append(uyari("Yüksek", "Zamanaşımı yaklaşıyor", kim, f"Zamanaşımı tarihi {d['zaman']:%d.%m.%Y} ({d['zaman_kalan']} gün kaldı). {bilgi}"))
        elif d["talep"] is None:
            d["durum"] = "Talep gönderilmedi"
            uy.append(uyari("Orta", "Rücu talebi gönderilmedi", kim, bilgi))
        elif d["bekleme"] is not None and d["bekleme"] > cevap and d["asama"] in ("Hazırlık", "Talep gönderildi"):
            d["durum"] = "Cevap gecikti"
            uy.append(uyari("Orta", "Cevap gecikti", kim, f"Talep {d['talep']:%d.%m.%Y} tarihinde gönderildi, {d['bekleme']} gündür cevap yok. {bilgi}"))
        elif d["asama"] == "Reddedildi":
            d["durum"] = "Ret, karar bekliyor"
            uy.append(uyari("Orta", "Reddedilen rücu için karar bekleniyor", kim, f"{bilgi} Tahkim, dava ya da kapatma kararı verilmeli."))
        elif d["asama"] == "Kabul edildi" and d["cevap"] and (bugun - d["cevap"]).days > tahsil:
            d["durum"] = "Tahsilat gecikti"
            uy.append(uyari("Orta", "Kabul edilen rücu tahsil edilmedi", kim, f"Kabul {d['cevap']:%d.%m.%Y} ({(bugun - d['cevap']).days} gün önce). {bilgi}"))
        else:
            d["durum"] = "Süreçte"
        if d["zaman"] is None:
            uy.append(uyari("Bilgi", "Zamanaşımı tarihi yazılmamış", kim, "Açık dosyada zamanaşımı tarihi yok; süre izlenemiyor."))

    def ozet(l: list[dict]) -> dict:
        acik, sureler = [d for d in l if d["acik"]], [d["cevap_sure"] for d in l if d["cevap_sure"] is not None]
        tutar, th = sum((d["tutar"] for d in l), Decimal(0)), sum((d["tahsil"] for d in l), Decimal(0))
        return {"dosya": len(l), "tutar": tutar, "tahsil": th, "oran": _oran(th, tutar), "acik": len(acik), "kalan": sum((d["kalan"] for d in acik), Decimal(0)),
                "vazgecilen": sum((d["kalan"] for d in l if not d["acik"]), Decimal(0)), "cevapsiz": sum(1 for d in acik if d["bekleme"] is not None), "ort_cevap": yuvarla(Decimal(sum(sureler)) / len(sureler)) if sureler else None}

    def kirilim(alan: str, sira=None) -> list[dict]:
        grup = {}
        for d in dosyalar:
            grup.setdefault(katla(d[alan]), []).append(d)
        s = [{"ad": g[0][alan], **ozet(g)} for g in grup.values()]
        return sorted(s, key=lambda x: (sira.index(x["ad"]) if x["ad"] in sira else len(sira), katla(x["ad"]))) if sira else sorted(s, key=lambda x: (-x["kalan"], -x["tutar"], katla(x["ad"])))

    return {"dosyalar": dosyalar, "acik": sorted((d for d in dosyalar if d["acik"]), key=lambda d: (DURUM.index(d["durum"]), -d["kalan"], katla(d["no"]))), "ozet": ozet(dosyalar),
            "asamalar": kirilim("asama", ASAMA_SIRA), "karsilar": kirilim("karsi"), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, bugun: date) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Açık Dosyalar"
    _baslik(ws, ["Durum", "Rücu No", "Hasar Dosya No", "Karşı Taraf", "Aşama", "Rücu Tutarı", "Tahsil Edilen", "Kalan", "Talep Tarihi", "Cevap Tarihi", "Cevap Bekleme (gün)", "Zamanaşımı Tarihi", "Zamanaşımına Kalan (gün)", "Sorumlu",
                 "Not (siz doldurun)"], (22, 11, 16, 28, 18, 15, 15, 15, 12, 12, 10, 12, 12, 18, 36))
    for d in r["acik"]:
        _satir(ws, [d["durum"], d["no"], d["hasar"], d["karsi"], d["asama"], d["tutar"], d["tahsil"] or None, d["kalan"], d["talep"], d["cevap"], d["bekleme"], d["zaman"], d["zaman_kalan"], d["sorumlu"], None], para_s=(6, 7, 8), tarih_s=(9, 10, 12),
               dolgu={1: DURUM_RENK[d["durum"]]})
    ws = wb.create_sheet("Özet")
    o = r["ozet"]
    _baslik(ws, ["Gösterge", "Değer"], (44, 18))
    for etiket, deger, para_mi in (("Rücu dosyası", o["dosya"], False), ("Toplam rücu tutarı", o["tutar"], True), ("Tahsil edilen", o["tahsil"], True), ("Tahsilat oranı %", o["oran"], True), ("Açık dosya", o["acik"], False),
                                   ("Açık dosyalarda kalan tutar", o["kalan"], True), ("Kapatılan dosyalarda tahsil edilemeyen", o["vazgecilen"], True), ("Cevap bekleyen dosya", o["cevapsiz"], False),
                                   ("Ortalama cevap süresi (gün)", o["ort_cevap"], True)):
        _satir(ws, [etiket, deger], para_s=(2,) if para_mi else ())
    ws.append([])
    _baslik(ws, ["Açık dosyaların durumu", "Dosya", "Kalan Tutar"], dondur="A2")
    for ad in DURUM:
        l = [d for d in r["acik"] if d["durum"] == ad]
        _satir(ws, [ad, len(l), sum((d["kalan"] for d in l), Decimal(0))], para_s=(3,), dolgu={1: DURUM_RENK[ad] if l else None})
    ws.column_dimensions["C"].width = 18
    ws.append([])
    ws.append([f"Durum tarihi: {bugun:%d.%m.%Y}"])
    for baslik, alan, ad in (("Aşamalar", "asamalar", "Aşama"), ("Karşı Taraflar", "karsilar", "Karşı Taraf")):
        ws = wb.create_sheet(baslik)
        _baslik(ws, [ad, "Dosya", "Rücu Tutarı", "Tahsil Edilen", "Tahsilat %", "Açık Dosya", "Kalan Tutar", "Cevap Bekleyen", "Ort. Cevap Süresi (gün)"], (30, 8, 16, 16, 11, 9, 16, 10, 12))
        for s in r[alan]:
            _satir(ws, [s["ad"], s["dosya"], s["tutar"], s["tahsil"], s["oran"], s["acik"] or None, s["kalan"] or None, s["cevapsiz"] or None, s["ort_cevap"]], para_s=(3, 4, 5, 7, 9))
    ws = wb.create_sheet("Tüm Dosyalar")
    _baslik(ws, ["Rücu No", "Hasar Dosya No", "Karşı Taraf", "Aşama", "Durum", "Rücu Tutarı", "Tahsil Edilen", "Tahsilat %", "Talep Tarihi", "Cevap Tarihi", "Cevap Süresi (gün)", "Sorumlu"], (11, 16, 28, 18, 22, 15, 15, 10, 12, 12, 10, 18))
    for d in r["dosyalar"]:
        _satir(ws, [d["no"], d["hasar"], d["karsi"], d["asama"], d["durum"], d["tutar"], d["tahsil"], d["oran"], d["talep"], d["cevap"], d["cevap_sure"], d["sorumlu"]], para_s=(6, 7, 8), tarih_s=(9, 10), dolgu={5: DURUM_RENK[d["durum"]]})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, gun: int = 90, cevap: int = 30, tahsil: int = 30) -> dict:
    dosyalar, uy = oku(yol)
    r = degerlendir(dosyalar, bugun, gun, cevap, tahsil)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, bugun)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Rücu dosyalarını talep, cevap, tahkim/dava ve tahsilat aşamalarında izler.")
    p.add_argument("--dosyalar", type=Path, help="Rücu No, Hasar Dosya No, Karşı Taraf, Rücu Tutarı, Talep Tarihi, Cevap Tarihi, Aşama, Tahsil Edilen, Zamanaşımı Tarihi, Sorumlu")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--gun", type=int, default=90, help="Zamanaşımına kaç gün kala uyarılsın (varsayılan 90)")
    p.add_argument("--cevap", type=int, default=30, help="Talepten sonra kaç günden fazla cevapsızlık gecikme sayılsın (varsayılan 30)")
    p.add_argument("--tahsil", type=int, default=30, help="Kabulden sonra kaç günden fazla tahsilatsızlık gecikme sayılsın (varsayılan 30)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "rucu_tahsilat_takibi.xlsx")
    a = p.parse_args(argv)
    ornek = a.dosyalar is None
    yol = ORNEK / "rucu_dosyalari.csv" if ornek else a.dosyalar
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    if bugun is None or a.gun < 0 or a.cevap < 0 or a.tahsil < 0:
        print("[X] --tarih GG.AA.YYYY; --gun, --cevap ve --tahsil eksi olamaz")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, bugun, a.gun, a.cevap, a.tahsil)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {o['dosya']} rücu dosyası · talep {o['tutar']:,.2f} · tahsil {o['tahsil']:,.2f} (%{o['oran']}) · açık {o['acik']} dosya, kalan {o['kalan']:,.2f}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
