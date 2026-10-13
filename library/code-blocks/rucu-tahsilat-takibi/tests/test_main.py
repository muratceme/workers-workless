import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)


def g(n):
    return None if n is None else BUGUN + timedelta(days=n)


def d(no, tutar=1000, talep=-10, cevap=None, asama="Talep gönderildi", tahsil=0, zaman=400, karsi="K"):
    """talep, cevap, zaman: bugünden kaç gün sonra (eksi = geçmiş; None = boş)."""
    return {"no": no, "hasar": "", "karsi": karsi, "tutar": D(tutar), "talep": g(talep), "cevap": g(cevap), "asama": asama, "tahsil": D(tahsil), "zaman": g(zaman), "sorumlu": ""}


class DegerlendirTesti(unittest.TestCase):
    def test_durumlar(self):
        l = [d("geçmiş", zaman=-1), d("z90", zaman=90), d("z91", zaman=91), d("z0", zaman=0), d("talepsiz", talep=None, asama="Hazırlık"), d("c30", talep=-30), d("c31", talep=-31), d("tahkimde", talep=-200, asama="Tahkim"),
             d("ret", talep=-50, cevap=-40, asama="Reddedildi"), d("kabul30", talep=-50, cevap=-30, asama="Kabul edildi"), d("kabul31", talep=-50, cevap=-31, asama="Kabul edildi", tahsil=400),
             d("tam", talep=-50, cevap=-40, asama="Kabul edildi", tahsil=1000), d("kısmi", talep=-50, cevap=-40, asama="Tahsil edildi", tahsil=700), d("kapalı", talep=-50, cevap=-40, asama="Kapatıldı", zaman=-5),
             d("zamansız", zaman=None), d("fazla", asama="Tahsil edildi", tahsil=1100)]
        r = main.degerlendir(l, BUGUN)
        self.assertEqual({x["no"]: (x["durum"], x["kalan"], x["acik"]) for x in r["dosyalar"]},
                         {"geçmiş": ("Zamanaşımı geçmiş", D(1000), True), "z90": ("Zamanaşımı yaklaşıyor", D(1000), True), "z91": ("Süreçte", D(1000), True), "z0": ("Zamanaşımı yaklaşıyor", D(1000), True),
                          "talepsiz": ("Talep gönderilmedi", D(1000), True), "c30": ("Süreçte", D(1000), True), "c31": ("Cevap gecikti", D(1000), True), "tahkimde": ("Süreçte", D(1000), True),
                          "ret": ("Ret, karar bekliyor", D(1000), True), "kabul30": ("Süreçte", D(1000), True), "kabul31": ("Tahsilat gecikti", D(600), True), "tam": ("Tahsil edildi", D(0), False),
                          "kısmi": ("Kısmen tahsil edildi", D(300), False), "kapalı": ("Kapatıldı", D(1000), False), "zamansız": ("Süreçte", D(1000), True), "fazla": ("Tahsil edildi", D(0), False)})
        self.assertEqual([x["no"] for x in r["acik"]], ["geçmiş", "z0", "z90", "talepsiz", "c31", "ret", "kabul31", "c30", "kabul30", "tahkimde", "z91", "zamansız"])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         sorted([("Yüksek", "Zamanaşımı tarihi geçmiş", "geçmiş · K"), ("Yüksek", "Zamanaşımı yaklaşıyor", "z90 · K"), ("Yüksek", "Zamanaşımı yaklaşıyor", "z0 · K"), ("Orta", "Rücu talebi gönderilmedi", "talepsiz · K"),
                                 ("Orta", "Cevap gecikti", "c31 · K"), ("Orta", "Reddedilen rücu için karar bekleniyor", "ret · K"), ("Orta", "Kabul edilen rücu tahsil edilmedi", "kabul31 · K"),
                                 ("Bilgi", "Zamanaşımı tarihi yazılmamış", "zamansız · K"), ("Orta", "Tahsilat rücu tutarından büyük", "fazla · K")]))
        o = r["ozet"]
        self.assertEqual((o["dosya"], o["tutar"], o["tahsil"], o["oran"], o["acik"], o["kalan"], o["vazgecilen"], o["cevapsiz"], o["ort_cevap"]), (16, D(16000), D(3200), D("20.00"), 12, D(11600), D(1300), 8, D("13.17")))
        p = lambda **kw: main.degerlendir([d("a", talep=-20, zaman=100)], BUGUN, **kw)["dosyalar"][0]["durum"]   # noqa: E731
        self.assertEqual((p(), p(gun=100), p(cevap=19), p(cevap=20)), ("Süreçte", "Zamanaşımı yaklaşıyor", "Cevap gecikti", "Süreçte"))
        self.assertEqual(main.degerlendir([d("a", talep=-60, cevap=-20, asama="Kabul edildi")], BUGUN, tahsil=19)["dosyalar"][0]["durum"], "Tahsilat gecikti")
        with self.assertRaises(ValueError):
            main.degerlendir([], BUGUN)

    def test_kirilim(self):
        l = [d("1", 1000, -40, -30, "Kabul edildi", 400, karsi="A"), d("2", 2000, -40, None, karsi="a"), d("3", 500, -40, -20, "Tahsil edildi", 500, karsi="B"), d("4", 300, None, None, "Özel aşama", karsi="B"), d("5", 800, -90, -60, "Dava / İcra", karsi="B")]
        r = main.degerlendir(l, BUGUN)
        k = lambda alan: [(s["ad"], s["dosya"], s["tutar"], s["tahsil"], s["oran"], s["acik"], s["kalan"], s["cevapsiz"], s["ort_cevap"]) for s in r[alan]]   # noqa: E731
        self.assertEqual(k("karsilar"), [("A", 2, D(3000), D(400), D("13.33"), 2, D(2600), 1, D("10.00")), ("B", 3, D(1600), D(500), D("31.25"), 2, D(1100), 0, D("25.00"))])
        self.assertEqual([x[:3] for x in k("asamalar")], [("Talep gönderildi", 1, D(2000)), ("Kabul edildi", 1, D(1000)), ("Dava / İcra", 1, D(800)), ("Tahsil edildi", 1, D(500)), ("Özel aşama", 1, D(300))])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "rucu_dosyalari.csv", cikti, BUGUN)
            self.assertEqual([(x["no"], x["durum"], x["kalan"]) for x in r["acik"]],
                             [("R-007", "Zamanaşımı geçmiş", D(31000)), ("R-004", "Zamanaşımı yaklaşıyor", D(230000)), ("R-005", "Talep gönderilmedi", D(48000)), ("R-002", "Cevap gecikti", D(140000)), ("R-006", "Tahsilat gecikti", D(56000)),
                              ("R-009", "Süreçte", D(120000)), ("R-003", "Süreçte", D(62000))])
            self.assertEqual({x["no"]: (x["bekleme"], x["zaman_kalan"], x["cevap_sure"]) for x in r["dosyalar"] if x["no"] in ("R-002", "R-004", "R-007", "R-001")},
                             {"R-002": (72, 220, None), "R-004": (None, 49, 23), "R-007": (None, -41, 26), "R-001": (None, 154, 22)})
            o = r["ozet"]
            self.assertEqual((o["dosya"], o["tutar"], o["tahsil"], o["oran"], o["acik"], o["kalan"], o["vazgecilen"], o["cevapsiz"]), (9, D(866000), D(145000), D("16.74"), 7, D(687000), D(34000), 2))
            self.assertEqual([(s["ad"], s["dosya"], s["kalan"], s["cevapsiz"], s["ort_cevap"]) for s in r["karsilar"]],
                             [("Örnek Sigorta A.Ş.", 3, D(350000), 0, D("23.00")), ("Deneme Sigorta A.Ş.", 2, D(196000), 1, D("18.00")), ("Numune Sigorta A.Ş.", 2, D(93000), 1, D("26.00")), ("Kurgu Nakliyat Ltd. Şti.", 2, D(48000), 0, D("23.00"))])
            self.assertEqual([s["ad"] for s in r["asamalar"]], ["Hazırlık", "Talep gönderildi", "Kabul edildi", "Reddedildi", "Tahkim", "Tahsil edildi", "Kapatıldı"])
            self.assertEqual([(u["onem"], u["tur"], u["kim"][:5]) for u in r["uyarilar"]],
                             [("Yüksek", "Zamanaşımı tarihi geçmiş", "R-007"), ("Yüksek", "Zamanaşımı yaklaşıyor", "R-004"), ("Orta", "Cevap gecikti", "R-002"), ("Orta", "Kabul edilen rücu tahsil edilmedi", "R-006"),
                              ("Orta", "Rücu talebi gönderilmedi", "R-005")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Açık Dosyalar", "Özet", "Aşamalar", "Karşı Taraflar", "Tüm Dosyalar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "r.csv"
            y.write_text("No;Dosya No;Muhatap;Talep Tutarı;Gönderim Tarihi;Yanıt Tarihi;Durum;Tahsilat;Zamanaşımı\nR1;H1;A;1.250,50;01.09.2026;x;KABUL;250,50;01.01.2027\nR2;;;500;;;İcra;;\nR3;;;500;;;Bilinmeyen;;\nr1;;;500;;;;;\n"
                         "R4;;;0;;;;;\nR5;;;100;;;;x;\n;;;100;;;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["hasar"], x["karsi"], x["tutar"], x["talep"], x["cevap"], x["asama"], x["tahsil"], x["zaman"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("R1", "H1", "A", D("1250.50"), date(2026, 9, 1), None, "Kabul edildi", D("250.50"), date(2027, 1, 1)), ("R2", "", "(karşı taraf yok)", D(500), None, None, "Dava / İcra", D(0), None),
                               ("R3", "", "(karşı taraf yok)", D(500), None, None, "Bilinmeyen", D(0), None)],
                              [("Tarih okunamadı", "R1"), ("Aşama tanınmadı", "R3"), ("Dosya okunamadı", "r1"), ("Dosya okunamadı", "R4"), ("Dosya okunamadı", "R5")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--dosyalar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "12.10.2026", "--gun", "180", "--cevap", "15", "--tahsil", "45"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--gun", "-1"], ["--cevap", "-1"], ["--tahsil", "-1"])], [2] * 4)
            self.assertEqual(main.main(["--dosyalar", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Rücu No;Karşı Taraf\nA;B\n", "Rücu No;Rücu Tutarı\n;100\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
