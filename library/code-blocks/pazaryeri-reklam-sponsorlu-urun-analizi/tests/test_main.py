import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def s(urun, harcama, satis, tiklama=100, gosterim=10000, siparis=0, marj=None, kampanya="K"):
    return {"urun": urun, "kampanya": kampanya, "marj": None if marj is None else D(marj), "gosterim": D(gosterim), "tiklama": D(tiklama), "harcama": D(harcama), "satis": D(satis), "siparis": D(siparis)}


def tek(*a, **kw):
    ayar = {k: kw.pop(k) for k in ("hedef_acos", "min_tiklama") if k in kw}
    return main.analiz([s("X", *a, **kw)], **ayar)["urunler"][0]


class AnalizTesti(unittest.TestCase):
    def test_gostergeler(self):
        x = tek(1500, 12000, 500, 50000, 40, 25)
        self.assertEqual((x["ctr"], x["cpc"], x["donusum"], x["acos"], x["roas"], x["basabas"], x["kar"], x["sinif"]),
                         (D("1.00"), D("3.00"), D("8.00"), D("12.50"), D("8.00"), D(25), D("1500.00"), "Kârlı"))
        x = tek(0, 0, 0, 0)
        self.assertEqual((x["ctr"], x["cpc"], x["donusum"], x["acos"], x["roas"], x["sinif"]), (None, None, None, None, None, "Harcama yok"))

    def test_siniflar(self):
        self.assertEqual([tek(h, 1000, marj=20)["sinif"] for h in (200, "200.10")], ["Kârlı", "Kârsız"])                 # ACOS tam marja eşitse kârlı sayılır
        self.assertEqual([tek(50, 0, t, marj=20)["sinif"] for t in (19, 20)], ["Veri az", "Satışsız harcama"])
        self.assertEqual(tek(50, 0, 5, marj=20, min_tiklama=D(5))["sinif"], "Satışsız harcama")
        self.assertEqual((tek(300, 1000)["sinif"], tek(300, 1000, hedef_acos=D(25))["sinif"], tek(300, 1000, hedef_acos=D(30))["sinif"]), ("Marj bilinmiyor", "Kârsız", "Kârlı"))
        x = tek(300, 1000, hedef_acos=D(25))
        self.assertEqual((x["basabas"], x["kar"]), (D(25), None))                                                        # hedef ACOS kâr hesabında kullanılmaz
        self.assertEqual(tek(1800, 6000, marj=20)["kar"], D("-600.00"))

    def test_toplama_ve_siralama(self):
        l = [s("T", 1500, 12000, 500, 50000, 40, 25, "Genel"), s("t", 300, 3000, 100, 10000, 10, 25, "Marka"), s("J", 1800, 6000, 200, 20000, 8, 20, "Genel"), s("Ş", 400, 0, 80, marj=30), s("Ç", 500, 2500)]
        r = main.analiz(l)
        t = next(u for u in r["urunler"] if u["urun"] == "T")
        self.assertEqual((t["harcama"], t["satis"], t["tiklama"], t["donusum"], t["acos"], t["roas"], t["kar"]), (D(1800), D(15000), D(600), D("8.33"), D("12.00"), D("8.33"), D("1950.00")))
        self.assertEqual([(u["urun"], u["sinif"]) for u in r["urunler"]], [("Ş", "Satışsız harcama"), ("J", "Kârsız"), ("Ç", "Marj bilinmiyor"), ("T", "Kârlı")])
        g = r["genel"]
        self.assertEqual((g["harcama"], g["satis"], g["acos"], g["roas"], g["karsiz_harcama"], g["karsiz_pay"], g["kar"]), (D(4500), D(23500), D("19.15"), D("5.22"), D(2200), D("48.89"), D("950.00")))
        self.assertEqual([(k["kampanya"], k["harcama"], k["acos"]) for k in r["kampanyalar"]], [("Genel", D(3300), D("18.33")), ("K", D(900), D("36.00")), ("Marka", D(300), D("10.00"))])
        self.assertEqual([u["tur"] for u in main.analiz([s("A", 1, 10, marj=20), s("A", 1, 10, marj=30)])["uyarilar"]], ["Ürünün marjı satırlarda farklı"])
        with self.assertRaises(ValueError):
            main.analiz([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "r.xlsx"
            r = main.calistir(ORNEK / "reklam_raporu.csv", cikti)
            self.assertEqual([(u["urun"].split()[0], u["sinif"], u["acos"], u["kar"]) for u in r["urunler"]],
                             [("Şapka", "Satışsız harcama", None, D("-400.00")), ("Slim", "Kârsız", D("30.00"), D("-600.00")), ("Omuz", "Marj bilinmiyor", D("20.00"), None),
                              ("Mevsimlik", "Veri az", None, D("-60.00")), ("Basic", "Kârlı", D("12.00"), D("1950.00"))])
            g = r["genel"]
            self.assertEqual((g["harcama"], g["satis"], g["acos"], g["roas"], g["karsiz_harcama"], g["karsiz_pay"]), (D(4560), D(23500), D("19.40"), D("5.15"), D(2200), D("48.25")))
            self.assertEqual([(k["kampanya"], k["harcama"], k["acos"]) for k in r["kampanyalar"]], [("Genel", D(3700), D("20.56")), ("Marka", D(860), D("15.64"))])
            self.assertEqual(([u["tur"] for u in r["uyarilar"]], load_workbook(cikti).sheetnames), (["Marjı olmayan ürünler var"], ["Ürün Bazında", "Kampanya Bazında", "Uyarılar"]))

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "r.csv"
            y.write_text("Ürün Adı;Gösterim;Tıklama;Reklam Harcaması;Satış;Sipariş;Marj\nA;100;200;1.250,50;5.000;çok;150\nB;1;1;yok;1;1;10\nC;;;10;0;;%20\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual([(x["urun"], x["harcama"], x["satis"], x["siparis"], x["marj"], x["kampanya"]) for x in l], [("A", D("1250.50"), D(5000), D(0), None, "Belirtilmemiş"), ("C", D(10), D(0), D(0), D(20), "Belirtilmemiş")])
            self.assertEqual(sorted(u["tur"] for u in uy), ["Değer okunamadı", "Harcama veya satış okunamadı", "Marj geçersiz", "Tıklama gösterimden fazla"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--reklam", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--hedef-acos", "22,5", "--min-tiklama", "10"]), 0)
            self.assertEqual(main.main(arg + ["--hedef-acos", "120"]), 2)
            self.assertEqual(main.main(arg + ["--min-tiklama", "az"]), 2)
            self.assertEqual(main.main(["--reklam", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Ürün;Harcama;Reklam Satışı\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
