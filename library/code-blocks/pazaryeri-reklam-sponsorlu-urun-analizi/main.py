"""
Pazaryeri Reklam (Sponsorlu Ürün) Analizi — Gerisi Bende kod bloğu
Performans Pazarlama › Performans Pazarlama Uzmanı

Pazaryeri reklam raporlarından ürün bazında reklam maliyeti/satış oranını hesaplayıp kârsız reklamları işaretler.
  Tıklama oranı (CTR) % = tıklama ÷ gösterim × 100          Tıklama başı maliyet (CPC) = harcama ÷ tıklama
  Dönüşüm %            = sipariş ÷ tıklama × 100            ROAS = reklam satışı ÷ harcama
  ACOS %               = harcama ÷ reklam satışı × 100      (reklam maliyetinin reklam satışına oranı)
  Başabaş ACOS         = ürünün reklam öncesi kâr marjı %   (dosyadaki "Marj %"; yoksa --hedef-acos)
  Reklam sonrası kâr   = reklam satışı × marj % − harcama
  Sınıf: Harcama yok · Veri az (satış yok ve tıklama < --min-tiklama 20) · Satışsız harcama · Kârsız (ACOS > başabaş ACOS) · Kârlı · Marj bilinmiyor
Marj; ürün maliyeti, pazaryeri komisyonu, kargo ve iade payı düşüldükten sonraki orandır ve sizin verinizdir. Fiyat ve harcamalar aynı KDV esasında olmalıdır.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"urun": ("urun", "urun adi", "sku", "stok kodu", "barkod"), "kampanya": ("kampanya", "kampanya adi", "reklam grubu"), "gosterim": ("gosterim", "goruntulenme", "gosterim sayisi"),
         "tiklama": ("tiklama", "tiklama sayisi", "tik"), "harcama": ("harcama", "reklam harcamasi", "maliyet", "butce harcamasi"), "satis": ("reklam satisi", "satis", "ciro", "reklam cirosu", "satis tutari"),
         "siparis": ("siparis", "siparis sayisi", "satis adedi", "adet"), "marj": ("marj", "marj %", "kar marji", "marj yuzde")}
SAYI = ("gosterim", "tiklama", "harcama", "satis", "siparis")
SIRA = ("Satışsız harcama", "Kârsız", "Marj bilinmiyor", "Veri az", "Kârlı", "Harcama yok")
Z = Decimal(0)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("urun", "harcama", "satis")):
        u = metin(r["urun"])
        if not u:
            continue
        kim = f"{u} (satır {r['_satir']})"
        v = {a: para(r[a]) for a in SAYI}
        if v["harcama"] is None or v["satis"] is None or v["harcama"] < 0 or v["satis"] < 0:
            uy.append(uyari("Yüksek", "Harcama veya satış okunamadı", kim, f"Harcama '{metin(r['harcama'])}', reklam satışı '{metin(r['satis'])}'. Satır alınmadı."))
            continue
        for a in ("gosterim", "tiklama", "siparis"):
            if v[a] is None and metin(r[a]) or (v[a] is not None and v[a] < 0):
                uy.append(uyari("Orta", "Değer okunamadı", kim, f"{SUTUN[a][0]} '{metin(r[a])}'. 0 sayıldı."))
                v[a] = None
            v[a] = v[a] or Z
        marj = para(r["marj"])
        if marj is not None and not 0 < marj < 100:
            uy.append(uyari("Orta", "Marj geçersiz", kim, f"Marj '{metin(r['marj'])}' 0 ile 100 arasında olmalı. Boş sayıldı."))
            marj = None
        if v["tiklama"] > v["gosterim"] > 0:
            uy.append(uyari("Orta", "Tıklama gösterimden fazla", kim, f"Tıklama {v['tiklama']:f}, gösterim {v['gosterim']:f}. Sütunları kontrol edin."))
        l.append({"urun": u, "kampanya": metin(r["kampanya"]) or "Belirtilmemiş", "marj": marj} | v)
    return l, uy


def olcut(x: dict) -> dict:
    return {"ctr": yuvarla(x["tiklama"] / x["gosterim"] * 100) if x["gosterim"] else None, "cpc": yuvarla(x["harcama"] / x["tiklama"]) if x["tiklama"] else None,
            "donusum": yuvarla(x["siparis"] / x["tiklama"] * 100) if x["tiklama"] else None, "acos": yuvarla(x["harcama"] / x["satis"] * 100) if x["satis"] else None,
            "roas": yuvarla(x["satis"] / x["harcama"]) if x["harcama"] else None}


def analiz(satirlar: list[dict], hedef_acos: Decimal | None = None, min_tiklama: Decimal = Decimal(20)) -> dict:
    if not satirlar:
        raise ValueError("Analiz edilecek reklam satırı yok.")
    uy = []

    def grupla(alan: str) -> list[dict]:
        g = {}
        for s in satirlar:
            x = g.setdefault(katla(s[alan]), {alan: s[alan], "marjlar": []} | {a: Z for a in SAYI})
            for a in SAYI:
                x[a] += s[a]
            if s["marj"] is not None:
                x["marjlar"].append(s["marj"])
        for x in g.values():
            x |= olcut(x)
        return list(g.values())

    urunler = grupla("urun")
    for u in urunler:
        if len(set(u["marjlar"])) > 1:
            uy.append(uyari("Orta", "Ürünün marjı satırlarda farklı", u["urun"], f"Marj değerleri: {', '.join(f'{m:f}' for m in sorted(set(u['marjlar'])))}. İlk satırdaki kullanıldı."))
        u["marj"] = u["marjlar"][0] if u["marjlar"] else None
        u["basabas"] = u["marj"] if u["marj"] is not None else hedef_acos
        u["kar"] = yuvarla(u["satis"] * u["marj"] / 100 - u["harcama"]) if u["marj"] is not None else None
        if u["harcama"] == 0:
            u["sinif"] = "Harcama yok"
        elif u["satis"] == 0:
            u["sinif"] = "Veri az" if u["tiklama"] < min_tiklama else "Satışsız harcama"
        elif u["basabas"] is None:
            u["sinif"] = "Marj bilinmiyor"
        else:
            u["sinif"] = "Kârsız" if u["acos"] > u["basabas"] else "Kârlı"
    kampanyalar = grupla("kampanya")
    genel = {a: sum((s[a] for s in satirlar), Z) for a in SAYI}
    genel |= olcut(genel)
    karsiz = [u for u in urunler if u["sinif"] in ("Kârsız", "Satışsız harcama")]
    genel |= {"karsiz_harcama": sum((u["harcama"] for u in karsiz), Z), "kar": sum((u["kar"] for u in urunler if u["kar"] is not None), Z) if any(u["kar"] is not None for u in urunler) else None}
    genel["karsiz_pay"] = yuvarla(genel["karsiz_harcama"] / genel["harcama"] * 100) if genel["harcama"] else Z
    return {"urunler": sorted(urunler, key=lambda u: (SIRA.index(u["sinif"]), -u["harcama"], u["urun"])), "kampanyalar": sorted(kampanyalar, key=lambda k: (-k["harcama"], k["kampanya"])), "genel": genel,
            "say": {s: sum(u["sinif"] == s for u in urunler) for s in SIRA}, "ayar": (hedef_acos, min_tiklama), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    renk = {"Satışsız harcama": KIRMIZI, "Kârsız": KIRMIZI, "Marj bilinmiyor": SARI, "Veri az": MAVI, "Kârlı": YESIL}
    ha, mt = r["ayar"]
    bas = ["Gösterim", "Tıklama", "CTR %", "CPC", "Sipariş", "Dönüşüm %", "Harcama", "Reklam Satışı", "ACOS %", "ROAS"]
    deger = lambda x: [x["gosterim"], x["tiklama"], x["ctr"], x["cpc"], x["siparis"], x["donusum"], x["harcama"], x["satis"], x["acos"], x["roas"]]   # noqa: E731
    ws = wb.active
    ws.title = "Ürün Bazında"
    _baslik(ws, ["Sınıf", "Ürün", *bas, "Marj %", "Başabaş ACOS %", "Reklam Sonrası Kâr", "Aksiyon (siz doldurun)"], (18, 34, 11, 9, 8, 9, 9, 10, 13, 14, 9, 8, 9, 11, 14, 30))
    for u in r["urunler"]:
        _satir(ws, [u["sinif"], u["urun"], *deger(u), u["marj"], u["basabas"], u["kar"], ""], para_s=(5, 6, 8, 9, 10, 11, 12, 13, 14, 15),
               dolgu={1: renk.get(u["sinif"]), 15: (KIRMIZI if u["kar"] < 0 else None) if u["kar"] is not None else None})
    g = r["genel"]
    _satir(ws, ["Toplam", "", *deger(g), None, None, g["kar"], ""], para_s=(5, 6, 8, 9, 10, 11, 12, 15))
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.append([])
    ws.append([f"Kârsız ve satışsız reklamların harcaması {g['karsiz_harcama']:,.2f} (toplam harcamanın %{g['karsiz_pay']}). Başabaş ACOS = ürünün marjı"
               + (f"; marjı olmayan ürünlerde hedef ACOS %{ha.normalize():f}" if ha is not None else "; marjı olmayan ürünler sınıflandırılmadı")
               + f". Satışı olmayan ve tıklaması {mt.normalize():f} altındaki ürünler 'Veri az' sayıldı."])
    ws = wb.create_sheet("Kampanya Bazında")
    _baslik(ws, ["Kampanya", *bas], (30, 11, 9, 8, 9, 9, 10, 13, 14, 9, 8))
    for k in r["kampanyalar"]:
        _satir(ws, [k["kampanya"], *deger(k)], para_s=(4, 5, 7, 8, 9, 10, 11))
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, hedef_acos: Decimal | None = None, min_tiklama: Decimal = Decimal(20)) -> dict:
    satirlar, uy = oku(yol)
    r = analiz(satirlar, hedef_acos, min_tiklama)
    uy += r["uyarilar"]
    if r["say"]["Marj bilinmiyor"]:
        uy.append(uyari("Orta", "Marjı olmayan ürünler var", yol.name, f"{r['say']['Marj bilinmiyor']} ürün kârlı veya kârsız olarak sınıflandırılamadı. Dosyaya 'Marj %' sütunu ekleyin veya --hedef-acos verin."))
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Pazaryeri reklam raporundan ürün bazında ACOS hesaplar, kârsız reklamları işaretler.")
    p.add_argument("--reklam", type=Path, default=ORNEK / "reklam_raporu.csv", help="Ürün, Kampanya, Gösterim, Tıklama, Harcama, Reklam Satışı, Sipariş, Marj %%")
    p.add_argument("--hedef-acos", help="Marjı yazılmamış ürünler için başabaş ACOS %% (verilmezse bu ürünler sınıflandırılmaz)")
    p.add_argument("--min-tiklama", default="20", help="Satışı olmayan ürünü 'Satışsız harcama' saymak için en az tıklama (varsayılan 20)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "reklam_analizi.xlsx")
    a = p.parse_args(argv)
    ha, mt = para(a.hedef_acos) if a.hedef_acos else None, para(a.min_tiklama)
    if a.hedef_acos and (ha is None or not 0 < ha < 100) or mt is None or mt < 0:
        print("[X] --hedef-acos 0 ile 100 arasında; --min-tiklama sıfır veya pozitif sayı olmalı")
        return 2
    if not a.reklam.exists():
        print(f"[X] Dosya bulunamadı: {a.reklam}")
        return 1
    try:
        r = calistir(a.reklam, a.cikti, ha, mt)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    g, s = r["genel"], r["say"]
    print(f"[OK] {len(r['urunler'])} ürün · harcama {g['harcama']:,.2f} · reklam satışı {g['satis']:,.2f} · ACOS %{g['acos'] if g['acos'] is not None else '—'} · ROAS {g['roas'] if g['roas'] is not None else '—'}")
    print(f"[{'!' if g['karsiz_harcama'] else 'OK'}] {s['Kârsız']} kârsız, {s['Satışsız harcama']} satışsız harcama: {g['karsiz_harcama']:,.2f} (harcamanın %{g['karsiz_pay']})")
    for u in r["urunler"]:
        if u["sinif"] in ("Kârsız", "Satışsız harcama"):
            print(f"[!] {u['urun']}: {u['sinif']} · harcama {u['harcama']:,.2f}" + (f" · ACOS %{u['acos']} > başabaş %{u['basabas'].normalize():f}" if u["acos"] is not None else ""))
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
