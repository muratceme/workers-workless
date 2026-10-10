import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def calistir(tmp, **kw):
    return main.calistir(ORNEK / "fiyat_listesi.csv", kw.pop("kalem", ORNEK / "teklif_kalemleri.csv"), kw.pop("bilgi", ORNEK / "teklif_bilgisi.csv"), Path(tmp),
                         ORNEK / "iskonto_kurallari.csv", **kw)


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.s = calistir(cls.tmp.name)
        cls.k = {k.urun.kod: k for k in cls.s["kalemler"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_elle_hesaplanan_kalemler(self):
        p = self.k["PMP-100"]                                   # 850 EUR × 48,10 = 40.885; × 0,85 (bayi) × 0,95 (kampanya)
        self.assertEqual((p.liste_tl, p.net, p.tutar), (D("40885.0000"), D("33014.64"), D("132058.56")))
        v = self.k["VAN-050"]                                   # 60 adet: yalnız en yüksek miktar kademesi (%10), eski kampanya geçersiz
        self.assertEqual([a for a, _ in v.uygulanan], ["Bayi iskontosu", "Vana 50 adet ve üzeri"])
        self.assertEqual(v.net, D("1874.25"))                  # 2.450 × 0,85 × 0,90
        b = self.k["BRU-025"]                                   # 1.180 × 0,85 × 0,93 × 0,95 = 886,1505
        self.assertEqual(b.net, D("886.15"))
        self.assertLess(b.marj, 0)
        self.assertEqual(self.k["OTM-300"].liste_tl, D("21450.0000"))   # 520 USD × 41,25

    def test_toplamlar_ve_onay(self):
        s = self.s
        self.assertEqual(s["ara"], sum((k.tutar for k in s["kalemler"]), D(0)))
        self.assertEqual(s["toplam_kdv"], (s["ara"] * D("0.20")).quantize(D("0.01")))
        self.assertEqual(s["genel"], s["ara"] + s["toplam_kdv"])
        self.assertTrue(s["onay"])
        turler = {(u["tur"], u["kim"].split()[0]) for u in s["uyarilar"]}
        self.assertTrue({("Fiyat listesinde yok", "XYZ-999"), ("Düşük marj", "BRU-025"), ("Asgari sipariş altı", "VAN-080")} <= turler)

    def test_dosyalar(self):
        t = load_workbook(self.s["teklif"]).active
        icerik = " ".join(str(c.value) for r in t.iter_rows() for c in r if c.value is not None)
        self.assertIn("FİYAT TEKLİFİ", icerik)
        self.assertNotIn("Maliyet", icerik)                     # iç bilgi müşteri dosyasına girmez
        self.assertNotIn("arj", icerik)
        self.assertEqual(load_workbook(self.s["ic"]).sheetnames, ["Özet", "Hesap", "Uyarılar"])


class KuralTesti(unittest.TestCase):
    def test_eksik_kur(self):
        with tempfile.TemporaryDirectory() as tmp:
            bilgi = Path(tmp) / "b.csv"
            bilgi.write_text("Alan;Değer\nTeklif No;T1\nTarih;09.10.2026\nMüşteri Grubu;Proje\nKur EUR;48,10\n", encoding="utf-8")
            with self.assertRaises(ValueError):                 # OTM-300 USD; kur yok
                calistir(tmp, bilgi=bilgi)
            kalem = Path(tmp) / "k.csv"
            kalem.write_text("Ürün Kodu;Miktar\nVAN-080;10\n", encoding="utf-8")
            s = calistir(tmp, bilgi=bilgi, kalem=kalem)
            self.assertEqual(s["kalemler"][0].net, D("3588.00"))          # proje %8
            self.assertFalse(s["onay"])

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", tmp]), 0)
            self.assertEqual(main.main(["--teklif", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
