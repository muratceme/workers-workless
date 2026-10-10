"""
Teklif Hazırlama — Workers / Workless kod bloğu
Satış › Satış Temsilcisi

Fiyat listesi, iskonto kuralları ve ürün seçiminden müşteriye gönderilecek teklif dosyasını hazırlar:
  - Liste fiyatı teklif dövizine çevrilir (teklif bilgisindeki kurlarla).
  - İskonto kuralları (müşteri grubu, kategori, ürün, miktar kademesi, kampanya; geçerlilik tarihleriyle) satıra
    uygulanır. Aynı "Grup" içindeki kurallardan yalnız en yüksek oran geçerlidir; farklı gruplar zincirleme uygulanır:
    net = liste × (1 − i1) × (1 − i2) ... Satırdaki ek (manuel) iskonto en son zincire eklenir.
  - Net birim fiyat kuruşa yuvarlanır; tutar = miktar × net birim fiyat. KDV oran bazında ara toplam üzerinden.
  - İç kontrol (müşteriye gitmez): maliyet varsa marj = (net − maliyet) / net; marj < --min-marj veya toplam iskonto
    > --max-iskonto ise yönetici onayı gerekir. Asgari sipariş miktarı altı ve fiyat listesinde olmayan ürün uyarısı.
Çıktı: müşteriye gönderilecek teklif (.xlsx; başlık, satırlar, KDV dökümü, şartlar) ve ayrı iç kontrol dosyası (hesap
adımları, marj, onay gereksinimi). İnternete bağlanmaz.

Kullanım:
    python main.py                                                  # örnek: 8 kalemlik bayi teklifi
    python main.py --fiyat-listesi fiyat.xlsx --iskontolar iskonto.xlsx --teklif kalemler.xlsx --bilgi teklif_bilgisi.csv
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
K2 = Decimal("0.01")
SIFIR = Decimal(0)
BIR = Decimal(1)

FIYAT_SUTUNLARI = {"kod": ("urun kodu", "kod", "stok kodu"), "ad": ("urun", "urun adi", "aciklama"), "kategori": ("kategori", "urun grubu"),
                   "birim": ("birim", "olcu birimi"), "fiyat": ("liste fiyati", "fiyat", "birim fiyat"), "doviz": ("doviz", "para birimi"),
                   "kdv": ("kdv orani", "kdv"), "maliyet": ("maliyet", "birim maliyet"), "moq": ("asgari siparis", "minimum siparis")}
ISKONTO_SUTUNLARI = {"kural": ("kural", "kural adi"), "tur": ("tur", "kosul turu"), "kosul": ("kosul", "deger"), "min": ("min miktar", "asgari miktar"),
                     "oran": ("iskonto", "iskonto orani", "oran"), "grup": ("grup", "iskonto grubu"), "bas": ("baslangic", "baslangic tarihi"),
                     "bit": ("bitis", "bitis tarihi")}
KALEM_SUTUNLARI = {"kod": ("urun kodu", "kod", "stok kodu"), "miktar": ("miktar", "adet"), "ek": ("ek iskonto", "ozel iskonto", "manuel iskonto"),
                   "not": ("not", "aciklama")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("%", "").replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def oran(x) -> Decimal | None:
    v = para(x)
    return None if v is None else (v / 100 if v > 1 else v)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def yuzde(o: Decimal, basamak: int | None = None) -> str:
    v = o * 100
    yazi = f"{v:.{basamak}f}" if basamak is not None else f"{v.normalize():f}"
    return "%" + yazi.replace(".", ",")


def tutar_yaz(x: Decimal, doviz: str = "") -> str:
    s = f"{x.quantize(K2, ROUND_HALF_UP):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"{s} {doviz}".strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

def bilgi_oku(yol: Path) -> dict:
    """'Alan;Değer' satırları: Teklif No, Tarih, Müşteri, Müşteri Grubu, İlgili Kişi, Geçerlilik (gün), Döviz, Kur USD, Kur EUR,
    Ödeme, Teslim Şekli, Teslim Süresi, Hazırlayan, Firma, Notlar."""
    b = {"satirlar": {}}
    for r in tablo_oku(yol):
        if len(r) >= 2 and metin(r[0]):
            b["satirlar"][metin(r[0])] = r[1]
    s = {katla(k): v for k, v in b["satirlar"].items()}
    b["no"] = metin(s.get("teklif no")) or "TEKLİF"
    b["tarih"] = tarih(s.get("tarih")) or date.today()
    b["musteri"] = metin(s.get("musteri"))
    b["grup"] = metin(s.get("musteri grubu"))
    b["gecerlilik"] = int(para(s.get("gecerlilik gun")) or para(s.get("gecerlilik")) or 15)
    b["doviz"] = (metin(s.get("doviz")) or "TRY").upper().replace("TL", "TRY")
    b["kurlar"] = {"TRY": BIR}
    for k, v in s.items():
        m = re.fullmatch(r"kur ([a-z]{3})", k)
        if m and para(v):
            b["kurlar"][m[1].upper()] = para(v)
    for alan in ("ilgili kisi", "odeme", "teslim sekli", "teslim suresi", "hazirlayan", "firma", "notlar"):
        b[alan.replace(" ", "_")] = metin(s.get(alan))
    return b


@dataclass
class Urun:
    kod: str
    ad: str
    kategori: str
    birim: str
    fiyat: Decimal
    doviz: str
    kdv: Decimal
    maliyet: Decimal | None
    moq: Decimal | None


@dataclass
class Kalem:
    sira: int
    urun: Urun
    miktar: Decimal
    ek: Decimal | None
    not_: str
    liste_tl: Decimal = SIFIR          # teklif dövizinde liste birim fiyatı
    uygulanan: list[tuple[str, Decimal]] = field(default_factory=list)
    net: Decimal = SIFIR
    tutar: Decimal = SIFIR
    toplam_iskonto: Decimal = SIFIR
    marj: Decimal | None = None
    maliyet: Decimal | None = None
    notlar: list[str] = field(default_factory=list)


def fiyat_oku(yol: Path) -> dict[str, Urun]:
    u = {}
    for r in kayitlar(yol, FIYAT_SUTUNLARI, ("kod", "fiyat")):
        kod, f = metin(r.get("kod")).upper(), para(r.get("fiyat"))
        if kod and f is not None:
            u[kod] = Urun(kod, metin(r.get("ad")), metin(r.get("kategori")), metin(r.get("birim")) or "Adet", f, (metin(r.get("doviz")) or "TRY").upper().replace("TL", "TRY"),
                          oran(r.get("kdv")) if metin(r.get("kdv")) else Decimal("0.20"), para(r.get("maliyet")), para(r.get("moq")))
    return u


def iskonto_oku(yol: Path | None) -> list[dict]:
    if not yol:
        return []
    lst = []
    for r in kayitlar(yol, ISKONTO_SUTUNLARI, ("tur", "oran")):
        o = oran(r.get("oran"))
        if o is None:
            continue
        lst.append({"kural": metin(r.get("kural")) or metin(r.get("tur")), "tur": katla(r.get("tur")), "kosul": metin(r.get("kosul")), "min": para(r.get("min")),
                    "oran": o, "grup": metin(r.get("grup")) or metin(r.get("kural")), "bas": tarih(r.get("bas")), "bit": tarih(r.get("bit"))})
    return lst


# ----------------------------------------------------------------------------
# Hesap
# ----------------------------------------------------------------------------

def kural_uyar_mi(k: dict, kalem: Kalem, musteri_grubu: str, gun: date) -> bool:
    if (k["bas"] and gun < k["bas"]) or (k["bit"] and gun > k["bit"]):
        return False
    if k["min"] is not None and kalem.miktar < k["min"]:
        return False
    t, kosul = k["tur"], katla(k["kosul"])
    if t.startswith("musteri"):
        return kosul == katla(musteri_grubu)
    if t.startswith("kategori"):
        return kosul == katla(kalem.urun.kategori)
    if t.startswith("urun"):
        return kosul == katla(kalem.urun.kod)
    if t.startswith("miktar") or t.startswith("kampanya") or t.startswith("genel"):
        return not kosul or kosul in (katla(kalem.urun.kategori), katla(kalem.urun.kod), "tumu", "hepsi")
    return False


def kalem_hesapla(kalem: Kalem, kurallar: list[dict], bilgi: dict) -> None:
    u = kalem.urun
    kur_urun, kur_teklif = bilgi["kurlar"].get(u.doviz), bilgi["kurlar"].get(bilgi["doviz"])
    if kur_urun is None or kur_teklif is None:
        raise ValueError(f"{u.kod}: {u.doviz} → {bilgi['doviz']} çevrimi için teklif bilgisinde 'Kur {u.doviz if kur_urun is None else bilgi['doviz']}' yok")
    kalem.liste_tl = u.fiyat * kur_urun / kur_teklif
    if u.doviz != bilgi["doviz"]:
        kalem.notlar.append(f"{tutar_yaz(u.fiyat, u.doviz)} × kur {str(kur_urun).replace('.', ',')}" + ("" if kur_teklif == BIR else f" / {str(kur_teklif).replace('.', ',')}"))
    en_iyi = {}
    for k in kurallar:
        if kural_uyar_mi(k, kalem, bilgi["grup"], bilgi["tarih"]):
            g = katla(k["grup"])
            if g not in en_iyi or k["oran"] > en_iyi[g]["oran"]:
                en_iyi[g] = k
    kalem.uygulanan = [(k["kural"], k["oran"]) for k in en_iyi.values()]
    if kalem.ek:
        kalem.uygulanan.append(("Ek iskonto", kalem.ek))
    carpan = BIR
    for _, o in kalem.uygulanan:
        carpan *= (1 - o)
    kalem.net = (kalem.liste_tl * carpan).quantize(K2, ROUND_HALF_UP)
    kalem.tutar = (kalem.net * kalem.miktar).quantize(K2, ROUND_HALF_UP)
    kalem.toplam_iskonto = 1 - carpan
    if u.maliyet is not None:
        kalem.maliyet = u.maliyet * kur_urun / kur_teklif
        kalem.marj = (kalem.net - kalem.maliyet) / kalem.net if kalem.net else None


def hesapla(kalemler: list[Kalem], kurallar: list[dict], bilgi: dict, min_marj: Decimal, max_iskonto: Decimal) -> dict:
    uy = []
    for k in kalemler:
        kalem_hesapla(k, kurallar, bilgi)
        ad = f"{k.urun.kod} {k.urun.ad}"
        if k.marj is not None and k.marj < min_marj:
            uy.append({"onem": "Yüksek" if k.marj < 0 else "Orta", "tur": "Düşük marj", "kim": ad,
                       "aciklama": f"Net {tutar_yaz(k.net, bilgi['doviz'])}, maliyet {tutar_yaz(k.maliyet, bilgi['doviz'])} → marj {yuzde(k.marj, 1)} "
                                   f"(alt sınır {yuzde(min_marj, 0)}); yönetici onayı gerekir"})
        if k.toplam_iskonto > max_iskonto:
            uy.append({"onem": "Orta", "tur": "İskonto sınırı aşıldı", "kim": ad,
                       "aciklama": f"Toplam iskonto {yuzde(k.toplam_iskonto, 1)} (sınır {yuzde(max_iskonto, 0)}): "
                                   + " × ".join(f"{a} {yuzde(o)}" for a, o in k.uygulanan)})
        if k.urun.moq and k.miktar < k.urun.moq:
            uy.append({"onem": "Bilgi", "tur": "Asgari sipariş altı", "kim": ad, "aciklama": f"Miktar {k.miktar:g}, asgari sipariş {k.urun.moq:g} {k.urun.birim}"})
    kdv = defaultdict(Decimal)
    for k in kalemler:
        kdv[k.urun.kdv] += k.tutar
    kdv_dokum = [(o, matrah, (matrah * o).quantize(K2, ROUND_HALF_UP)) for o, matrah in sorted(kdv.items())]
    ara = sum((k.tutar for k in kalemler), SIFIR)
    liste_toplam = sum(((k.liste_tl * k.miktar).quantize(K2, ROUND_HALF_UP) for k in kalemler), SIFIR)
    toplam_kdv = sum((x[2] for x in kdv_dokum), SIFIR)
    maliyetli = [k for k in kalemler if k.maliyet is not None]
    toplam_maliyet = sum((k.maliyet * k.miktar for k in maliyetli), SIFIR)
    maliyetli_ciro = sum((k.tutar for k in maliyetli), SIFIR)
    genel_marj = (maliyetli_ciro - toplam_maliyet) / maliyetli_ciro if maliyetli_ciro else None
    onay = any(u["tur"] in ("Düşük marj", "İskonto sınırı aşıldı") for u in uy)
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    uy.sort(key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))
    return {"kalemler": kalemler, "kdv": kdv_dokum, "ara": ara, "liste_toplam": liste_toplam, "iskonto_tutari": liste_toplam - ara, "toplam_kdv": toplam_kdv,
            "genel": ara + toplam_kdv, "genel_marj": genel_marj, "onay": onay, "uyarilar": uy, "bilgi": bilgi,
            "gecerlilik_tarihi": bilgi["tarih"] + timedelta(days=bilgi["gecerlilik"])}


# ----------------------------------------------------------------------------
# Çıktılar
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
INCE = Side(style="thin", color="BBBBBB")
CERCEVE = Border(top=INCE, bottom=INCE, left=INCE, right=INCE)
PF = "#,##0.00"


def teklif_yaz(yol: Path, s: dict) -> None:
    b = s["bilgi"]
    dv = b["doviz"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Teklif"
    for j, w in enumerate((5, 12, 40, 9, 7, 13, 9, 13, 15), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.append([b.get("firma") or ""])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["FİYAT TEKLİFİ"])
    ws["A2"].font = Font(bold=True, size=12)
    ws.append([])
    for etiket, deger in (("Teklif No", b["no"]), ("Tarih", b["tarih"].strftime("%d.%m.%Y")), ("Sayın", b["musteri"]), ("İlgili Kişi", b.get("ilgili_kisi")),
                          ("Geçerlilik", f"{s['gecerlilik_tarihi']:%d.%m.%Y} tarihine kadar ({b['gecerlilik']} gün)"), ("Para Birimi", dv)):
        if deger:
            ws.append([etiket, None, deger])
            ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.append([])
    ws.append(["Sıra", "Ürün Kodu", "Ürün / Hizmet", "Miktar", "Birim", "Liste Fiyatı", "İskonto", "Net Birim Fiyat", "Tutar"])
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for k in s["kalemler"]:
        ws.append([k.sira, k.urun.kod, k.urun.ad, float(k.miktar), k.urun.birim, float(k.liste_tl.quantize(K2, ROUND_HALF_UP)), float(k.toplam_iskonto),
                   float(k.net), float(k.tutar)])
        for j in range(1, 10):
            ws.cell(ws.max_row, j).border = CERCEVE
        for j in (6, 8, 9):
            ws.cell(ws.max_row, j).number_format = PF
        ws.cell(ws.max_row, 7).number_format = "0.0%"
        ws.cell(ws.max_row, 3).alignment = Alignment(wrap_text=True)
    ws.append([])
    for etiket, v in [("Liste toplamı", s["liste_toplam"]), ("İskonto", -s["iskonto_tutari"]), ("Ara toplam (KDV hariç)", s["ara"])] + \
                     [(f"KDV {yuzde(o)} (matrah {tutar_yaz(m)})", k) for o, m, k in s["kdv"]] + [(f"GENEL TOPLAM ({dv})", s["genel"])]:
        ws.append([None, None, None, None, None, None, etiket, None, float(v)])
        ws.cell(ws.max_row, 9).number_format = PF
        ws.cell(ws.max_row, 7).alignment = Alignment(horizontal="left")
        if etiket.startswith("GENEL"):
            ws.cell(ws.max_row, 7).font = ws.cell(ws.max_row, 9).font = Font(bold=True)
    ws.append([])
    ws.append(["Şartlar"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for etiket, deger in (("Ödeme", b.get("odeme")), ("Teslim şekli", b.get("teslim_sekli")), ("Teslim süresi", b.get("teslim_suresi")),
                          ("Fiyatlar", "KDV hariçtir; KDV fatura tarihindeki oranla ayrıca eklenir." if dv == "TRY" else
                           f"{dv} cinsindendir; ödeme fatura / ödeme tarihindeki kurla yapılır, KDV ayrıca eklenir."),
                          ("Notlar", b.get("notlar"))):
        if deger:
            ws.append([etiket, None, deger])
            ws.cell(ws.max_row, 3).alignment = Alignment(wrap_text=True)
    if b.get("hazirlayan"):
        ws.append([])
        ws.append(["Hazırlayan", None, b["hazirlayan"]])
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(yol)


def ic_kontrol_yaz(yol: Path, s: dict) -> None:
    b = s["bilgi"]
    wb = Workbook()
    oz = wb.active
    oz.title = "Özet"
    for j, w in enumerate((38, 24), 1):
        oz.column_dimensions[get_column_letter(j)].width = w
    for etiket, v in [("Teklif No", b["no"]), ("Müşteri / grup", f"{b['musteri']} / {b['grup'] or '—'}"), ("Ara toplam", tutar_yaz(s["ara"], b["doviz"])),
                      ("Ortalama iskonto", yuzde(s["iskonto_tutari"] / s["liste_toplam"], 1) if s["liste_toplam"] else "—"),
                      ("Brüt marj (maliyeti bilinen kalemler)", "—" if s["genel_marj"] is None else yuzde(s["genel_marj"], 1)),
                      ("Yönetici onayı", "GEREKLİ" if s["onay"] else "Gerekmiyor")]:
        oz.append([etiket, v])
    oz.cell(oz.max_row, 2).fill = PatternFill("solid", fgColor="FDE2E1" if s["onay"] else "E3F4E1")
    oz.append([])
    oz.append(["Bu dosya şirket içi kullanım içindir; müşteriye gönderilmez."])

    hs = wb.create_sheet("Hesap")
    basliklar = ["Sıra", "Ürün Kodu", "Ürün", "Kategori", "Miktar", "Liste (ürün dövizi)", "Döviz", "Liste (teklif dövizi)", "Uygulanan İskontolar", "Toplam İskonto",
                 "Net Birim", "Tutar", "Birim Maliyet", "Marj", "Not"]
    hs.append(basliklar)
    for h in hs[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate((5, 11, 30, 14, 8, 12, 6, 12, 44, 9, 11, 13, 11, 8, 30), 1):
        hs.column_dimensions[get_column_letter(j)].width = w
    for k in s["kalemler"]:
        hs.append([k.sira, k.urun.kod, k.urun.ad, k.urun.kategori, float(k.miktar), float(k.urun.fiyat), k.urun.doviz, float(k.liste_tl),
                   " × ".join(f"{a} {yuzde(o)}" for a, o in k.uygulanan if o) or "—", float(k.toplam_iskonto), float(k.net), float(k.tutar),
                   None if k.maliyet is None else float(k.maliyet), None if k.marj is None else float(k.marj), "; ".join(k.notlar)])
        for j in (6, 8, 11, 12, 13):
            hs.cell(hs.max_row, j).number_format = PF
        hs.cell(hs.max_row, 10).number_format = hs.cell(hs.max_row, 14).number_format = "0.0%"
        if k.marj is not None and k.marj < s["min_marj"]:
            hs.cell(hs.max_row, 14).fill = PatternFill("solid", fgColor="FDE2E1")

    uy = wb.create_sheet("Uyarılar")
    uy.append(["Önem", "Tür", "Kalem", "Açıklama"])
    for h in uy[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for j, w in enumerate((9, 22, 34, 100), 1):
        uy.column_dimensions[get_column_letter(j)].width = w
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["kim"], u["aciklama"]])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(yol)


def calistir(fiyat_yolu: Path, kalem_yolu: Path, bilgi_yolu: Path, cikti_klasoru: Path, iskonto_yolu: Path | None = None, min_marj: Decimal = Decimal("0.15"),
             max_iskonto: Decimal = Decimal("0.25")) -> dict:
    urunler = fiyat_oku(fiyat_yolu)
    bilgi = bilgi_oku(bilgi_yolu)
    kalemler, uy = [], []
    for r in kayitlar(kalem_yolu, KALEM_SUTUNLARI, ("kod", "miktar")):
        kod, m = metin(r.get("kod")).upper(), para(r.get("miktar"))
        if not kod:
            continue
        if kod not in urunler:
            uy.append({"onem": "Yüksek", "tur": "Fiyat listesinde yok", "kim": kod, "aciklama": "Kalem teklife alınmadı; fiyat listesini güncelleyin"})
            continue
        if not m or m <= 0:
            uy.append({"onem": "Yüksek", "tur": "Geçersiz miktar", "kim": kod, "aciklama": f"Miktar '{metin(r.get('miktar'))}'; kalem alınmadı"})
            continue
        kalemler.append(Kalem(len(kalemler) + 1, urunler[kod], m, oran(r.get("ek")), metin(r.get("not"))))
    if not kalemler:
        raise ValueError("Teklife alınacak kalem yok")
    s = hesapla(kalemler, iskonto_oku(iskonto_yolu), bilgi, min_marj, max_iskonto)
    s["uyarilar"] = uy + s["uyarilar"]
    s["onay"] = s["onay"] or any(u["onem"] == "Yüksek" for u in uy)
    s["min_marj"] = min_marj
    cikti_klasoru.mkdir(parents=True, exist_ok=True)
    ad = re.sub(r"[^A-Za-z0-9_-]+", "_", katla(bilgi["no"]).replace(" ", "-")) or "teklif"
    s["teklif"] = cikti_klasoru / f"teklif_{ad}.xlsx"
    s["ic"] = cikti_klasoru / f"teklif_{ad}_ic_kontrol.xlsx"
    teklif_yaz(s["teklif"], s)
    ic_kontrol_yaz(s["ic"], s)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Fiyat listesi ve iskonto kurallarından müşteri teklif dosyası hazırlar.")
    p.add_argument("--fiyat-listesi", type=Path, default=ORNEK / "fiyat_listesi.csv",
                   help="Ürün Kodu, Ürün, Kategori, Birim, Liste Fiyatı, Döviz, KDV Oranı, Maliyet (iç), Asgari Sipariş")
    p.add_argument("--iskontolar", type=Path, help="Kural, Tür (Müşteri grubu / Kategori / Ürün / Miktar / Kampanya), Koşul, Min Miktar, İskonto, Grup, "
                                                   "Başlangıç, Bitiş")
    p.add_argument("--teklif", type=Path, default=ORNEK / "teklif_kalemleri.csv", help="Ürün Kodu, Miktar, Ek İskonto, Not")
    p.add_argument("--bilgi", type=Path, default=ORNEK / "teklif_bilgisi.csv", help="Alan;Değer: Teklif No, Tarih, Müşteri, Müşteri Grubu, Döviz, Kur USD...")
    p.add_argument("--min-marj", type=float, default=15, help="Bu marjın (%%) altında yönetici onayı (varsayılan 15)")
    p.add_argument("--max-iskonto", type=float, default=25, help="Toplam iskonto bu %%'yi aşarsa yönetici onayı (varsayılan 25)")
    p.add_argument("--cikti", type=Path, default=Path("cikti"), help="Çıktı klasörü")
    a = p.parse_args(argv)
    ornek = a.fiyat_listesi == ORNEK / "fiyat_listesi.csv"
    isk = a.iskontolar or (ORNEK / "iskonto_kurallari.csv" if ornek else None)
    for y in (a.fiyat_listesi, a.teklif, a.bilgi, isk):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        s = calistir(a.fiyat_listesi, a.teklif, a.bilgi, a.cikti, isk, Decimal(str(a.min_marj)) / 100, Decimal(str(a.max_iskonto)) / 100)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    b = s["bilgi"]
    print(f"[OK] {b['no']} · {b['musteri']} · {len(s['kalemler'])} kalem · ara toplam {tutar_yaz(s['ara'], b['doviz'])} · genel toplam "
          f"{tutar_yaz(s['genel'], b['doviz'])}")
    print(f"[{'!' if s['onay'] else 'OK'}] Yönetici onayı {'GEREKLİ' if s['onay'] else 'gerekmiyor'}"
          + ("" if s["genel_marj"] is None else f" · brüt marj {yuzde(s['genel_marj'], 1)}"))
    for u in s["uyarilar"]:
        if u["onem"] != "Bilgi":
            print(f"[!] {u['kim']} · {u['tur']}: {u['aciklama']}")
    print(f"[OK] Teklif: {s['teklif'].resolve()}")
    print(f"[OK] İç kontrol: {s['ic'].resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
