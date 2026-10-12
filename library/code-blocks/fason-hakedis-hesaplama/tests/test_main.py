import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def st(verilen=1025, saglam=1000, defolu=15, fiyat=40, malzeme=180, avans=0, diger=0, atolye="A"):
    return {"atolye": atolye, "siparis": "S", "model": "M", "kim": atolye, "verilen": D(verilen), "saglam": D(saglam), "defolu": D(defolu), "fiyat": D(fiyat),
            "malzeme": None if malzeme is None else D(malzeme), "avans": D(avans), "diger": D(diger)}


class HesapTesti(unittest.TestCase):
    def test_satir(self):
        s = main.hesapla([st()], tolerans=D(1))["satirlar"][0]
        self.assertEqual((s["eksik"], s["brut"], s["tol_adet"], s["tabi"], s["kesinti"], s["defolu_oran"], s["isaret"]), (D(10), D("40000.00"), D("10.25"), D("14.75"), D("2655.00"), D("1.46"), ""))
        s = main.hesapla([st()])["satirlar"][0]                                              # tolerans yok: 25 adet × 180
        self.assertEqual((s["tabi"], s["kesinti"]), (D(25), D("4500.00")))
        s = main.hesapla([st(822, 815, 5)], tolerans=D(1))["satirlar"][0]                     # 7 adet < 8,22 tolerans
        self.assertEqual((s["tabi"], s["kesinti"]), (D(0), D("0.00")))
        s = main.hesapla([st(500, 460, 25, 65, 320)], tolerans=D(1))["satirlar"][0]
        self.assertEqual((s["eksik"], s["tabi"], s["kesinti"], s["defolu_oran"], s["isaret"]), (D(15), D(35), D("11200.00"), D("5.00"), "Defolu oranı yüksek"))

    def test_atolye_net_odeme(self):
        r = main.hesapla([st(avans=10000), st(822, 815, 5)], tolerans=D(1), kdv=D(10), tevkifat=D("0.7"))
        a = r["atolye"][0]
        self.assertEqual((a["brut"], a["kdv"], a["tevkif"], a["kesinti"], a["avans"], a["net"]), (D("72600.00"), D("7260.00"), D("5082.00"), D("2655.00"), D(10000), D("62123.00")))
        a = main.hesapla([st(avans=10000)])["atolye"][0]                                      # KDV ve tolerans verilmezse
        self.assertEqual((a["kdv"], a["tevkif"], a["net"]), (D("0.00"), D("0.00"), D("25500.00")))

    def test_uyarilar(self):
        r = main.hesapla([st(300, 310, 0)])
        self.assertEqual((r["satirlar"][0]["eksik"], [u["tur"] for u in r["uyarilar"]]), (D(0), ["Teslim verilenden fazla"]))
        self.assertEqual([u["tur"] for u in main.hesapla([st(malzeme=None)])["uyarilar"]], ["Malzeme bedeli yok"])
        r = main.hesapla([st(100, 10, 0, 40, 180, avans=50000)])
        self.assertEqual((r["atolye"][0]["net"] < 0, [u["tur"] for u in r["uyarilar"]]), (True, ["Net ödeme eksi"]))
        with self.assertRaises(ValueError):
            main.hesapla([])
        self.assertEqual([main.kesir(x) for x in ("7/10", "0,5", "1/0", "x")], [D("0.7"), D("0.5"), None, None])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "f.xlsx"
            r = main.calistir(ORNEK / "fason_teslimler.csv", cikti, tolerans=D(1), kdv=D(10), tevkifat=D("0.7"))
            d = {a["atolye"][-1]: (a["brut"], a["kesinti"], a["diger"], a["avans"], a["kdv"], a["tevkif"], a["net"]) for a in r["atolye"]}
            self.assertEqual(d["A"], (D("72600.00"), D("2655.00"), D(0), D("10000.00"), D("7260.00"), D("5082.00"), D("62123.00")))
            self.assertEqual(d["B"], (D("46950.00"), D("11200.00"), D("1500.00"), D("5000.00"), D("4695.00"), D("3286.50"), D("30658.50")))
            self.assertEqual((r["toplam"]["net"], [u["tur"] for u in r["uyarilar"]]), (D("92781.50"), ["Teslim verilenden fazla"]))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Hakediş Özeti", "Satırlar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "t.csv"
            y.write_text("Atölye;Verilen Adet;Sağlam Adet;Defolu Adet;Birim Fiyat;Avans\nA;100;95;çok;10,50;-5\nB;yok;1;;1;\nC;10;10;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual([(s["atolye"], s["saglam"], s["defolu"], s["fiyat"], s["avans"], s["malzeme"]) for s in l], [("A", D(95), D(0), D("10.50"), D(0), None)])
            self.assertEqual(sorted(u["tur"] for u in uy), ["Adet veya fiyat okunamadı", "Adet veya fiyat okunamadı", "Değer okunamadı", "Değer okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--teslimler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tolerans", "2", "--kdv", "10", "--tevkifat", "7/10"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(arg + ["--tevkifat", "3/2"]), 2)
            self.assertEqual(main.main(arg + ["--kdv", "çok"]), 2)
            self.assertEqual(main.main(["--teslimler", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Atölye;Verilen Adet;Sağlam Adet;Birim Fiyat\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
