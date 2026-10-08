"""
Stok Sayım Fark Analizi — Workers / Workless kod bloğu
Lojistik ve Depo › Depo Sorumlusu

Sayım sonuçlarını sistem (ERP) stoğuyla karşılaştırır:
  - Lokasyon + stok kodu bazında fark (miktar ve tutar); sayım anında işlenmemiş hareketler (bekleyen irsaliye,
    üretim fişi vb.) verilirse sistem stoğu önce düzeltilir.
  - Ürün bazında net fark: bir lokasyondaki eksik başka lokasyondaki fazlayla kapanıyorsa "yer karışıklığı".
  - Olası neden ipuçları (kural tabanlı): sayılmamış lokasyon, fazla/eksik sıfır, rakam yer değiştirmesi,
    koli/adet karışıklığı, benzer kodlarla karşılıklı fark, negatif sistem stoğu, sistemde olmayan kalem.
  - Tolerans içi farklar (fire oranı) ayrılır; tutarı veya oranı yüksek farklar ikinci sayım listesine alınır.
  - Envanter doğruluğu (tolerans içindeki kalem oranı), lokasyon ve kategori özetleri,
    197 Sayım ve Tesellüm Noksanları / 397 Sayım ve Tesellüm Fazlaları için tutar özeti.
İnternete bağlanmaz.

Kullanım:
    python main.py                                         # örnek verilerle dener
    python main.py --sayim sayim.xlsx --sistem stok.xlsx
    python main.py --sayim sayim.xlsx --sistem stok.xlsx --bekleyen bekleyen.xlsx --tolerans 0.5 --esik 5000
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
TUM = "(tümü)"          # lokasyon bilgisi olmayan dosyalarda kullanılan ortak lokasyon


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def katla(s) -> str:
    """Başlık karşılaştırması için: Türkçe karakterleri ve ı/i farkını yok sayar ("IBAN" = "ıban" = "iban")."""
    s = kucuk(s)
    for a, b in zip("ıöüşğç", "iousgc"):
        s = s.replace(a, b)
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def sayi(x) -> Decimal:
    if x in (None, ""):
        return Decimal(0)
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("₺", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return Decimal(0)


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    satirlar = [r for r in satirlar if any(c not in (None, "") for c in r)]
    # Başlık satırı: "kod" içeren ilk satır (üstte başlık/açıklama satırları olabilir)
    for i, r in enumerate(satirlar[:10]):
        if any("kod" in katla(c) or katla(c) in {"sku", "barkod"} for c in r):
            return satirlar[i:]
    return satirlar


def bul(baslik: list, *adlar: str) -> int | None:
    b = [katla(x) for x in baslik]
    hedef = [katla(a) for a in adlar]
    return next((b.index(a) for a in hedef if a in b), None)


KOD = ("stok kodu", "malzeme kodu", "ürün kodu", "kod", "sku", "barkod")
LOK = ("lokasyon", "lokasyon kodu", "raf", "adres", "depo", "depo kodu", "yer")
AD = ("stok adı", "malzeme adı", "ürün adı", "açıklama", "ad")
MALIYET = ("birim maliyet", "birim fiyat", "maliyet", "ortalama maliyet")


def _hucre(r, i):
    return r[i] if i is not None and i < len(r) else None


def sistem_oku(yol: Path) -> tuple[dict, dict]:
    """Dönüş: ((lok, kod) → sistem miktarı, kod → {ad, maliyet, birim, kategori, koli})."""
    s = tablo_oku(yol)
    b = s[0]
    i_kod, i_lok, i_ad = bul(b, *KOD), bul(b, *LOK), bul(b, *AD)
    i_mik = bul(b, "sistem miktarı", "sistem stoğu", "kayıtlı miktar", "stok miktarı", "miktar", "bakiye", "mevcut")
    i_mal, i_bir = bul(b, *MALIYET), bul(b, "birim", "ölçü birimi")
    i_kat = bul(b, "kategori", "grup", "stok grubu", "ürün grubu")
    i_koli = bul(b, "koli içi", "koli içi adet", "paket içi", "ambalaj miktarı")
    if i_kod is None or i_mik is None:
        raise SystemExit(f"Sistem stoğu dosyasında Stok Kodu ve Sistem Miktarı (veya Miktar) gerekli. Başlıklar: {b}")
    miktar: dict[tuple[str, str], Decimal] = defaultdict(Decimal)
    kart: dict[str, dict] = {}
    for r in s[1:]:
        kod = str(_hucre(r, i_kod) or "").strip()
        if not kod:
            continue
        lok = (str(_hucre(r, i_lok) or "").strip() or TUM) if i_lok is not None else TUM
        miktar[(lok, kod)] += sayi(_hucre(r, i_mik))
        k = kart.setdefault(kod, {"ad": "", "maliyet": None, "birim": "", "kategori": "", "koli": None})
        if _hucre(r, i_ad):
            k["ad"] = str(_hucre(r, i_ad)).strip()
        if _hucre(r, i_mal) not in (None, ""):
            k["maliyet"] = sayi(_hucre(r, i_mal))
        if _hucre(r, i_bir):
            k["birim"] = str(_hucre(r, i_bir)).strip()
        if _hucre(r, i_kat):
            k["kategori"] = str(_hucre(r, i_kat)).strip()
        if _hucre(r, i_koli) not in (None, "") and sayi(_hucre(r, i_koli)) > 1:
            k["koli"] = sayi(_hucre(r, i_koli))
    return dict(miktar), kart


def sayim_oku(yol: Path) -> tuple[dict, dict, bool]:
    """Aynı lokasyon + kod için birden çok sayım satırı (farklı ekip/kart) toplanır.
    Dönüş: ((lok, kod) → sayılan, kod → ad (dosyada varsa), lokasyon sütunu var mı)."""
    s = tablo_oku(yol)
    b = s[0]
    i_kod, i_lok, i_ad = bul(b, *KOD), bul(b, *LOK), bul(b, *AD)
    i_mik = bul(b, "sayılan miktar", "sayım miktarı", "sayılan", "fiili miktar", "sayım", "miktar", "adet")
    if i_kod is None or i_mik is None:
        raise SystemExit(f"Sayım dosyasında Stok Kodu ve Sayılan Miktar gerekli. Başlıklar: {b}")
    sayilan: dict[tuple[str, str], Decimal] = defaultdict(Decimal)
    adlar = {}
    for r in s[1:]:
        kod = str(_hucre(r, i_kod) or "").strip()
        if not kod:
            continue
        lok = (str(_hucre(r, i_lok) or "").strip() or TUM) if i_lok is not None else TUM
        sayilan[(lok, kod)] += sayi(_hucre(r, i_mik))
        if _hucre(r, i_ad):
            adlar[kod] = str(_hucre(r, i_ad)).strip()
    return dict(sayilan), adlar, i_lok is not None


def bekleyen_oku(yol: Path | None, lokasyonlu: bool) -> dict:
    """Sayım anında sisteme işlenmemiş hareketler. Miktar işaretli (+ giriş, − çıkış) ya da Yön sütunuyla verilir."""
    if not yol:
        return {}
    s = tablo_oku(yol)
    b = s[0]
    i_kod, i_lok = bul(b, *KOD), bul(b, *LOK)
    i_mik = bul(b, "miktar", "adet")
    i_yon = bul(b, "yön", "hareket", "hareket türü", "giriş çıkış", "tür")
    if i_kod is None or i_mik is None:
        raise SystemExit(f"Bekleyen hareket dosyasında Stok Kodu ve Miktar gerekli. Başlıklar: {b}")
    sonuc: dict[tuple[str, str], Decimal] = defaultdict(Decimal)
    for r in s[1:]:
        kod = str(_hucre(r, i_kod) or "").strip()
        if not kod:
            continue
        lok = (str(_hucre(r, i_lok) or "").strip() or TUM) if (i_lok is not None and lokasyonlu) else TUM
        m = sayi(_hucre(r, i_mik))
        yon = katla(_hucre(r, i_yon))
        if yon.startswith(("cik", "sevk", "satis", "tuketim", "iade")):
            m = -abs(m)
        elif yon.startswith(("gir", "alis", "kabul", "uretim")):
            m = abs(m)
        sonuc[(lok, kod)] += m
    return dict(sonuc)


# ----------------------------------------------------------------------------
# Neden ipuçları
# ----------------------------------------------------------------------------

def _tamsayi(x: Decimal) -> int | None:
    return int(x) if x == x.to_integral_value() else None


def tek_satir_ipuclari(sistem: Decimal, sayilan: Decimal, koli: Decimal | None) -> list[str]:
    """Tek satıra bakarak çıkarılabilen olası nedenler. Kesin hüküm değil, kontrol önerisidir."""
    ip = []
    fark = sayilan - sistem
    if fark == 0:
        return ip
    if sistem < 0:
        ip.append("Sistem stoğu negatif: çıkış, girişten önce kaydedilmiş (kayıt sırası/eksik giriş fişi)")
    if sayilan == 0 and sistem > 0:
        ip.append("Hiç sayılmamış: lokasyon atlanmış veya ürün başka yere taşınmış olabilir")
    if sistem == 0 and sayilan > 0:
        ip.append("Sistemde bu lokasyonda stok yok: giriş/transfer fişi eksik olabilir")
    a, b = _tamsayi(sistem), _tamsayi(sayilan)
    if a and b and a > 0 and b > 0:
        buyuk, kucuk_ = max(a, b), min(a, b)
        if buyuk % kucuk_ == 0 and buyuk // kucuk_ in (10, 100, 1000):
            ip.append("Fazla/eksik sıfır: sayı giriş hatası olabilir")
        elif abs(a - b) % 9 == 0 and sorted(str(a)) == sorted(str(b)):
            ip.append("Rakamlar yer değiştirmiş olabilir (ör. 54 ↔ 45)")
    if koli and koli > 1 and sistem > 0 and sayilan > 0:
        if sayilan * koli == sistem or sistem * koli == sayilan:
            ip.append(f"Koli/adet karışıklığı: biri koli, diğeri adet olarak girilmiş olabilir (koli içi {koli:g})")
        elif fark % koli == 0:
            ip.append(f"Fark koli içi adedin ({koli:g}) tam katı: eksik/fazla sayılmış tam koli olabilir")
    return ip


def kod_benzer(a: str, b: str) -> bool:
    """Aynı ürün ailesinde varyant kodları (renk/beden/ebat): ortak kök uzun, tek bir parça farklı."""
    if a == b:
        return False
    pa, pb = re.split(r"[-_./ ]", a), re.split(r"[-_./ ]", b)
    if len(pa) == len(pb) and len(pa) > 1 and sum(x != y for x, y in zip(pa, pb)) == 1:
        return True
    return SequenceMatcher(None, a, b).ratio() >= 0.85


# ----------------------------------------------------------------------------
# Hesap
# ----------------------------------------------------------------------------

def calistir(sayim_yolu: Path, sistem_yolu: Path, cikti: Path, bekleyen_yolu: Path | None = None,
             tolerans: float = 0.0, esik_tutar: float = 5000.0, esik_oran: float = 10.0) -> dict:
    sayilan, sayim_adlari, lokasyonlu = sayim_oku(sayim_yolu)
    sistem, kart = sistem_oku(sistem_yolu)
    uyarilar = []
    if not lokasyonlu:                                  # sayımda lokasyon yoksa sistem de kod bazına indirilir
        tek = defaultdict(Decimal)
        for (_, kod), m in sistem.items():
            tek[(TUM, kod)] += m
        sistem = dict(tek)
    elif all(l == TUM for l, _ in sistem):              # sistemde lokasyon yoksa sayım kod bazına indirilir
        tek = defaultdict(Decimal)
        for (_, kod), m in sayilan.items():
            tek[(TUM, kod)] += m
        sayilan = dict(tek)
        uyarilar.append("Sistem stoğunda lokasyon yok: karşılaştırma stok kodu bazında yapıldı.")
    bekleyen = bekleyen_oku(bekleyen_yolu, lokasyonlu and not all(l == TUM for l, _ in sistem))

    satirlar = []
    for anahtar in sorted(set(sistem) | set(sayilan) | set(bekleyen)):
        lok, kod = anahtar
        k = kart.get(kod, {"ad": "", "maliyet": None, "birim": "", "kategori": "", "koli": None})
        s_ham = sistem.get(anahtar, Decimal(0))
        bek = bekleyen.get(anahtar, Decimal(0))
        s_duz = s_ham + bek
        say = sayilan.get(anahtar, Decimal(0))
        fark = say - s_duz
        maliyet = k["maliyet"]
        ip = tek_satir_ipuclari(s_duz, say, k["koli"])
        if kod not in kart:
            ip = ["Stok kartı sistemde yok: yanlış kod okunmuş/yazılmış veya kart açılmamış"]
        if bek:
            ip.append(f"Bekleyen hareketle düzeltildi ({bek:+g})")
        oran = float(abs(fark) / abs(s_duz) * 100) if s_duz else (100.0 if fark else 0.0)
        satirlar.append({
            "lok": lok, "kod": kod, "ad": k["ad"] or sayim_adlari.get(kod, ""), "kategori": k["kategori"] or "-",
            "birim": k["birim"], "sistem_ham": s_ham, "bekleyen": bek, "sistem": s_duz, "sayilan": say,
            "fark": fark, "oran": oran, "maliyet": maliyet,
            "tutar": (fark * maliyet) if maliyet is not None else None, "ipuclari": ip,
            "tolerans_ici": fark == 0 or (s_duz > 0 and oran <= tolerans),
        })
    eksik_maliyet = sorted({r["kod"] for r in satirlar if r["fark"] and r["maliyet"] is None})
    if eksik_maliyet:
        uyarilar.append(f"Birim maliyeti olmayan {len(eksik_maliyet)} kalemin tutar etkisi hesaplanamadı: "
                        + ", ".join(eksik_maliyet[:10]))

    # Ürün bazında: lokasyonlar arası karşılıklı farklar
    urun: dict[str, dict] = {}
    for r in satirlar:
        u = urun.setdefault(r["kod"], {"kod": r["kod"], "ad": r["ad"], "kategori": r["kategori"], "sistem": Decimal(0),
                                       "sayilan": Decimal(0), "eksik": Decimal(0), "fazla": Decimal(0),
                                       "maliyet": r["maliyet"], "lokasyonlar": []})
        u["sistem"] += r["sistem"]
        u["sayilan"] += r["sayilan"]
        if r["fark"] < 0:
            u["eksik"] += -r["fark"]
        elif r["fark"] > 0:
            u["fazla"] += r["fark"]
        if r["fark"]:
            u["lokasyonlar"].append(f"{r['lok']} {r['fark']:+g}")
    for u in urun.values():
        u["net"] = u["sayilan"] - u["sistem"]
        u["tutar"] = u["net"] * u["maliyet"] if u["maliyet"] is not None else None
        u["mahsup"] = min(u["eksik"], u["fazla"])        # lokasyonlar arasında kapanan miktar
    for r in satirlar:
        u = urun[r["kod"]]
        if r["fark"] and u["mahsup"] and r["lok"] != TUM:
            ip = "Başka lokasyonda ters yönde fark var: yer karışıklığı/transfer kaydı eksik olabilir"
            if u["net"] == 0:
                ip += " (ürün toplamında fark yok)"
            r["ipuclari"].insert(0, ip)

    # Benzer kodlar arasında karşılıklı fark (aynı lokasyonda, eşit miktar, ters işaret)
    farkli = [r for r in satirlar if r["fark"]]
    for i, r in enumerate(farkli):
        for t in farkli[i + 1:]:
            if r["lok"] == t["lok"] and r["fark"] == -t["fark"] and kod_benzer(r["kod"], t["kod"]):
                r["ipuclari"].append(f"{t['kod']} ile karşılıklı eşit fark: kod/varyant karışıklığı olabilir")
                t["ipuclari"].append(f"{r['kod']} ile karşılıklı eşit fark: kod/varyant karışıklığı olabilir")

    # İkinci sayım listesi
    for r in satirlar:
        nedenler = []
        if r["fark"] and not r["tolerans_ici"]:
            if r["tutar"] is not None and abs(r["tutar"]) >= Decimal(str(esik_tutar)):
                nedenler.append(f"tutar ≥ {esik_tutar:,.0f}".replace(",", "."))
            if r["oran"] >= esik_oran and r["sistem"]:
                nedenler.append(f"oran ≥ %{esik_oran:g}")
            if r["sayilan"] == 0 and r["sistem"] > 0:
                nedenler.append("hiç sayılmamış")
            if r["kod"] not in kart:
                nedenler.append("kartı yok")
        r["ikinci_sayim"] = ", ".join(nedenler)

    ozet = _ozet(satirlar, urun)
    _rapor(satirlar, urun, ozet, cikti, uyarilar, tolerans, esik_tutar, esik_oran)
    return {"satirlar": satirlar, "urun": urun, "ozet": ozet, "uyarilar": uyarilar}


def _ozet(satirlar, urun) -> dict:
    sifir = Decimal(0)
    eksik = sum((r["tutar"] for r in satirlar if r["tutar"] is not None and r["tutar"] < 0), sifir)
    fazla = sum((r["tutar"] for r in satirlar if r["tutar"] is not None and r["tutar"] > 0), sifir)
    u_eksik = sum((u["tutar"] for u in urun.values() if u["tutar"] is not None and u["tutar"] < 0), sifir)
    u_fazla = sum((u["tutar"] for u in urun.values() if u["tutar"] is not None and u["tutar"] > 0), sifir)
    sistem_degeri = sum((r["sistem"] * r["maliyet"] for r in satirlar if r["maliyet"] is not None), sifir)
    dogru = sum(1 for r in satirlar if r["tolerans_ici"])
    gruplar = {}
    for ad, alan in (("lokasyon", "lok"), ("kategori", "kategori")):
        g = {}
        for r in satirlar:
            x = g.setdefault(r[alan], {"satir": 0, "farkli": 0, "eksik": sifir, "fazla": sifir, "sistem_degeri": sifir})
            x["satir"] += 1
            x["farkli"] += 0 if r["tolerans_ici"] else 1
            if r["tutar"] is not None:
                x["eksik" if r["tutar"] < 0 else "fazla"] += r["tutar"]
            if r["maliyet"] is not None:
                x["sistem_degeri"] += r["sistem"] * r["maliyet"]
        gruplar[ad] = g
    return {"satir": len(satirlar), "dogru": dogru, "dogruluk": dogru / len(satirlar) * 100 if satirlar else 100.0,
            "eksik_tutar": eksik, "fazla_tutar": fazla, "net_tutar": eksik + fazla, "brut_tutar": fazla - eksik,
            "urun_eksik": u_eksik, "urun_fazla": u_fazla, "sistem_degeri": sistem_degeri,
            "ikinci_sayim": sum(1 for r in satirlar if r["ikinci_sayim"]), **gruplar}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
YESIL = PatternFill("solid", fgColor="E3F5E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
PARA = "#,##0.00"
MIK = "#,##0.###"


def _tl(x) -> str:
    return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _baslik(ws, satir: int = 1):
    for h in ws[satir]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI


def _f(x):
    return None if x is None else float(x)


def _rapor(satirlar, urun, ozet, cikti, uyarilar, tolerans, esik_tutar, esik_oran):
    wb = Workbook()
    o = wb.active
    o.title = "Özet"
    o.append(["Stok sayım fark analizi"])
    o["A1"].font = Font(bold=True, size=12)
    for s in [
        ("Karşılaştırılan satır (lokasyon × kalem)", ozet["satir"]),
        (f"Envanter doğruluğu (fark tolerans içinde, tolerans %{tolerans:g})", f"%{ozet['dogruluk']:.1f}".replace(".", ",")),
        ("Sistem stok değeri (düzeltilmiş)", _tl(ozet["sistem_degeri"])),
        ("Eksik tutarı (satır bazında)", _tl(ozet["eksik_tutar"])),
        ("Fazla tutarı (satır bazında)", _tl(ozet["fazla_tutar"])),
        ("Net fark tutarı", _tl(ozet["net_tutar"])),
        ("Brüt (mutlak) fark tutarı", _tl(ozet["brut_tutar"])),
        ("Ürün bazında net eksik (lokasyon mahsubu sonrası)", _tl(ozet["urun_eksik"])),
        ("Ürün bazında net fazla (lokasyon mahsubu sonrası)", _tl(ozet["urun_fazla"])),
        ("İkinci sayım önerilen satır", ozet["ikinci_sayim"]),
    ]:
        o.append(list(s))
    o.append([])
    o.append(["Lokasyon", "Satır", "Farklı satır", "Eksik tutarı", "Fazla tutarı", "Net", "Sistem değeri", "Brüt fark / değer"])
    _baslik(o, o.max_row)
    for lok, x in sorted(ozet["lokasyon"].items()):
        o.append([lok, x["satir"], x["farkli"], _f(x["eksik"]), _f(x["fazla"]), _f(x["eksik"] + x["fazla"]),
                  _f(x["sistem_degeri"]), _f((x["fazla"] - x["eksik"]) / x["sistem_degeri"]) if x["sistem_degeri"] else None])
        for c in (4, 5, 6, 7):
            o.cell(o.max_row, c).number_format = PARA
        o.cell(o.max_row, 8).number_format = "0.0%"
    o.append([])
    o.append(["Kategori", "Satır", "Farklı satır", "Eksik tutarı", "Fazla tutarı", "Net", "Sistem değeri", "Brüt fark / değer"])
    _baslik(o, o.max_row)
    for kat, x in sorted(ozet["kategori"].items()):
        o.append([kat, x["satir"], x["farkli"], _f(x["eksik"]), _f(x["fazla"]), _f(x["eksik"] + x["fazla"]),
                  _f(x["sistem_degeri"]), _f((x["fazla"] - x["eksik"]) / x["sistem_degeri"]) if x["sistem_degeri"] else None])
        for c in (4, 5, 6, 7):
            o.cell(o.max_row, c).number_format = PARA
        o.cell(o.max_row, 8).number_format = "0.0%"
    for u in uyarilar:
        o.append(["Uyarı", u])
    o.column_dimensions["A"].width = 52
    for c in "BCDEFGH":
        o.column_dimensions[c].width = 16

    f = wb.create_sheet("Farklar")
    f.append(["Lokasyon", "Stok Kodu", "Stok Adı", "Kategori", "Birim", "Sistem (ham)", "Bekleyen", "Sistem (düzeltilmiş)",
              "Sayılan", "Fark", "Fark %", "Birim Maliyet", "Fark Tutarı", "Tolerans İçi", "İkinci Sayım", "Olası Nedenler"])
    _baslik(f)
    for r in sorted(satirlar, key=lambda r: (r["tolerans_ici"], -abs(r["tutar"] or 0), r["lok"], r["kod"])):
        f.append([r["lok"], r["kod"], r["ad"], r["kategori"], r["birim"], _f(r["sistem_ham"]), _f(r["bekleyen"]) or None,
                  _f(r["sistem"]), _f(r["sayilan"]), _f(r["fark"]), r["oran"] / 100 if r["sistem"] else None,
                  _f(r["maliyet"]), _f(r["tutar"]), "Evet" if r["tolerans_ici"] else "Hayır", r["ikinci_sayim"],
                  " · ".join(r["ipuclari"])])
        n = f.max_row
        for c in (6, 7, 8, 9, 10):
            f.cell(n, c).number_format = MIK
        f.cell(n, 11).number_format = "0.0%"
        for c in (12, 13):
            f.cell(n, c).number_format = PARA
        if not r["tolerans_ici"]:
            f.cell(n, 10).fill = KIRMIZI if r["fark"] < 0 else SARI
        else:
            f.cell(n, 10).fill = YESIL
    for j, w in enumerate((12, 14, 28, 14, 7, 12, 10, 13, 10, 10, 8, 12, 14, 9, 22, 100), 1):
        f.column_dimensions[get_column_letter(j)].width = w
    f.freeze_panes = "C2"
    f.auto_filter.ref = f.dimensions

    u_ws = wb.create_sheet("Ürün Bazında")
    u_ws.append(["Stok Kodu", "Stok Adı", "Kategori", "Sistem", "Sayılan", "Net Fark", "Lokasyonlar Arası Mahsup",
                 "Birim Maliyet", "Net Fark Tutarı", "Farklı Lokasyonlar"])
    _baslik(u_ws)
    for u in sorted(urun.values(), key=lambda u: (-abs(u["tutar"] or 0), u["kod"])):
        if not (u["eksik"] or u["fazla"]):
            continue
        u_ws.append([u["kod"], u["ad"], u["kategori"], _f(u["sistem"]), _f(u["sayilan"]), _f(u["net"]), _f(u["mahsup"]) or None,
                     _f(u["maliyet"]), _f(u["tutar"]), ", ".join(u["lokasyonlar"])])
        for c in (8, 9):
            u_ws.cell(u_ws.max_row, c).number_format = PARA
    for j, w in enumerate((14, 28, 14, 10, 10, 10, 12, 12, 14, 50), 1):
        u_ws.column_dimensions[get_column_letter(j)].width = w
    u_ws.freeze_panes = "C2"

    i2 = wb.create_sheet("İkinci Sayım")
    i2.append(["Lokasyon", "Stok Kodu", "Stok Adı", "Sistem (düzeltilmiş)", "İlk Sayım", "Gerekçe", "İkinci Sayım", "Sayan", "Not"])
    _baslik(i2)
    for r in sorted((r for r in satirlar if r["ikinci_sayim"]), key=lambda r: (r["lok"], r["kod"])):
        i2.append([r["lok"], r["kod"], r["ad"], _f(r["sistem"]), _f(r["sayilan"]), r["ikinci_sayim"], None, None, None])
    for j, w in enumerate((12, 14, 28, 14, 10, 30, 13, 14, 30), 1):
        i2.column_dimensions[get_column_letter(j)].width = w
    i2.append([])
    i2.append(["Not: İkinci sayımı ilk sayımı yapmayan bir ekip, sistem miktarını görmeden (kör sayım) yapmalıdır."])

    m = wb.create_sheet("Muhasebe")
    m.append(["Sayım farklarının kayda alınması için tutar özeti (öneri — mali müşavirinizle kesinleştirin)"])
    m["A1"].font = Font(bold=True)
    m.append(["Hesap", "Tutar", "Açıklama"])
    _baslik(m, 2)
    m.append(["197 Sayım ve Tesellüm Noksanları", _f(-ozet["urun_eksik"]),
              "Ürün bazında net eksik (lokasyonlar arası mahsup sonrası) · karşılığı ilgili stok hesabı (150/151/152/153)"])
    m.append(["397 Sayım ve Tesellüm Fazlaları", _f(ozet["urun_fazla"]),
              "Ürün bazında net fazla · karşılığı ilgili stok hesabı"])
    for c in (3, 4):
        m.cell(c, 2).number_format = PARA
    for s in [
        "Bu hesaplar geçicidir: farkın nedeni araştırılıp belirlendikten sonra kapatılır (ör. kayıt hatası düzeltilir; "
        "nedeni bulunamayan eksik ve fazlalar dönem sonunda olağandışı gider/gelir hesaplarına aktarılır).",
        "Nedeni belgelenemeyen eksikler vergi incelemesinde kayıt dışı satış şüphesi doğurabilir; fire, zayi ve "
        "hırsızlık gibi durumlar tutanak ve belgeyle desteklenmelidir.",
        "Zayi olan mallara ait KDV'nin indirim hakkı ve normal fire oranlarının gider yazılması özel kurallara tabidir "
        "(KDV Kanunu md. 30). Uygulama için mali müşavirinize danışın.",
    ]:
        m.append(["Not", s])
    m.column_dimensions["A"].width = 36
    m.column_dimensions["B"].width = 16
    m.column_dimensions["C"].width = 120

    b = wb.create_sheet("Bilgi")
    for s in [["Fark", "Sayılan − sistem (düzeltilmiş). Negatif = eksik, pozitif = fazla"],
              ["Sistem (düzeltilmiş)", "Sistem miktarı + sayım anında işlenmemiş hareketler (bekleyen dosyası: + giriş, − çıkış)"],
              ["Tolerans içi", f"Fark sıfır veya |fark| / sistem ≤ %{tolerans:g}"],
              ["Envanter doğruluğu", "Tolerans içindeki satırların tüm satırlara oranı"],
              ["İkinci sayım", f"Tolerans dışı ve: |tutar| ≥ {_tl(esik_tutar)} TL, oran ≥ %{esik_oran:g}, hiç sayılmamış veya kartı olmayan satırlar"],
              ["Ürün bazında", "Aynı ürünün lokasyonlardaki farkları toplanır; eksik ve fazlanın küçüğü lokasyonlar arası mahsuptur"],
              ["Olası nedenler", "Kural tabanlı ipuçlarıdır, kesin tespit değildir; her biri yerinde kontrol edilmelidir"]]:
        b.append(s)
    b.column_dimensions["A"].width = 22
    b.column_dimensions["B"].width = 130
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ornek = BURASI / "ornek_veri"
    ap = argparse.ArgumentParser(description="Sayım sonuçlarını sistem stoğuyla karşılaştırır, farkları ve olası nedenleri raporlar.")
    ap.add_argument("--sayim", type=Path, default=ornek / "sayim.csv", help="Sayım (.xlsx/.csv): Lokasyon, Stok Kodu, Sayılan Miktar")
    ap.add_argument("--sistem", type=Path, default=ornek / "sistem_stogu.csv",
                    help="Sistem stoğu (.xlsx/.csv): Lokasyon, Stok Kodu, Stok Adı, Sistem Miktarı, Birim Maliyet [, Birim, Kategori, Koli İçi]")
    ap.add_argument("--bekleyen", type=Path, help="Sayım anında işlenmemiş hareketler (.xlsx/.csv): Lokasyon, Stok Kodu, Miktar [, Yön]")
    ap.add_argument("--tolerans", type=float, default=0.0, help="Kabul edilebilir fark oranı, %% (varsayılan 0)")
    ap.add_argument("--esik", type=float, default=5000.0, help="İkinci sayım için tutar eşiği, TL (varsayılan 5000)")
    ap.add_argument("--esik-oran", type=float, default=10.0, help="İkinci sayım için oran eşiği, %% (varsayılan 10)")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "sayim_fark_analizi.xlsx")
    a = ap.parse_args(argv)
    if a.sayim == ornek / "sayim.csv" and a.sistem == ornek / "sistem_stogu.csv" and not a.bekleyen:
        a.bekleyen = ornek / "bekleyen_hareketler.csv"
    s = calistir(a.sayim, a.sistem, a.cikti, a.bekleyen, a.tolerans, a.esik, a.esik_oran)
    o = s["ozet"]
    print(f"[OK] {o['satir']} satır · doğruluk %{o['dogruluk']:.1f} · eksik {_tl(o['eksik_tutar'])} · "
          f"fazla {_tl(o['fazla_tutar'])} · net {_tl(o['net_tutar'])} TL · ikinci sayım {o['ikinci_sayim']} satır")
    for u in s["uyarilar"]:
        print(f"[!] {u}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
