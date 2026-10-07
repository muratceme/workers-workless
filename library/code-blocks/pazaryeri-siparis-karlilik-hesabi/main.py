"""
Pazaryeri Sipariş Kârlılık Hesabı — Workers / Workless kod bloğu
E-ticaret › Pazaryeri Yönetimi › Pazaryeri Uzmanı

Pazaryeri sipariş dökümünü (Trendyol, Hepsiburada, Amazon, N11 vb.) ürün maliyet listesiyle birleştirip
sipariş satırı bazında net kârı hesaplar: satış KDV'si, pazaryeri komisyonu, kargo, platform hizmet bedeli,
ürün maliyeti ve iade kayıpları. Zararına satılan ürünleri ve başabaş satış fiyatını gösterir; pazaryerinin
ödeyeceği tutarı (e-ticaret stopajı %1 dahil) ve KDV etkisini ayrıca verir. İnternete bağlanmaz.

Kurallar kanal bazında kanallar.json dosyasından okunur (oranları kendi satıcı panelinizden güncelleyin):
  - komisyon_kdv_dahil: true → komisyon, KDV dahil satış tutarı × oran olarak hesaplanır ve bu tutar KDV'yi
    içerir (Trendyol Satıcı Bilgi Merkezi: "Komisyon faturaları KDV dahil rakamlardan hesaplanarak ... ekstra KDV
    ödemesi yapmanız gerekmez"). false → komisyon KDV hariç tutar × oran, üzerine KDV eklenir.
  - E-ticaret stopajı: aracı hizmet sağlayıcı, KDV hariç satış bedeli üzerinden %1 keser (9284 sayılı CBK,
    01.01.2025). Gelir/kurumlar vergisinden mahsup edildiği için kârı değil nakit akışını etkiler.

Kullanım:
    python main.py                                          # örnek siparişlerle dener
    python main.py --siparisler siparisler.xlsx --maliyetler urun_maliyetleri.xlsx --kanallar kanallar.json
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from collections import OrderedDict, defaultdict
from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
KURUS = Decimal("0.01")
SIFIR = Decimal(0)

KANAL_VARSAYILAN = {
    "komisyon_orani": "0",            # kategoriye özel oran sipariş dosyasında veya kategoriler'de verilebilir
    "kategoriler": {},                # {"kategori adı": oran}
    "komisyon_kdv_dahil": True,
    "hizmet_kdv_orani": "0.20",       # komisyon, kargo ve hizmet bedelinin KDV oranı
    "kargo_bedeli": "0",              # sipariş başına, KDV dahil (sipariş dosyasında Kargo sütunu varsa o kullanılır)
    "desi_tarifesi": {},              # {"üst desi": KDV dahil ücret} — maliyet listesinde Desi varsa
    "hizmet_bedeli": "0",             # sipariş başına platform hizmet bedeli, KDV dahil
    "stopaj_orani": "0.01",
    "iade_kargo_bedeli": None,        # iadede dönüş kargosu (KDV dahil); verilmezse gidiş kargosu kadar
    "iadede_hizmet_bedeli_kalir": True,
}


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def para(x) -> Decimal:
    if x in (None, ""):
        return SIFIR
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("₺", "").replace("%", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return SIFIR


def oran(x) -> Decimal:
    """%20, 20, 0,20 ve 0.20 hepsi 0.20 olarak okunur."""
    d = para(x)
    return d / 100 if d > 1 else d


def yuvarla(x: Decimal) -> Decimal:
    return x.quantize(KURUS, rounding=ROUND_HALF_UP)


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ayirici = csv.Sniffer().sniff(metin.splitlines()[0], delimiters=",;\t").delimiter
        satirlar = list(csv.reader(metin.splitlines(), delimiter=ayirici))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def sutunlar(baslik: list, alanlar: dict) -> dict:
    b = [kucuk(x) for x in baslik]
    return {alan: next((i for i, x in enumerate(b) if x in adlar), None) for alan, adlar in alanlar.items()}


SIPARIS_ALANLARI = {
    "siparis": ("sipariş no", "sipariş numarası", "siparis no", "order id", "paket no"),
    "tarih": ("tarih", "sipariş tarihi"),
    "kanal": ("kanal", "pazaryeri", "mağaza"),
    "kod": ("stok kodu", "barkod", "ürün kodu", "sku", "model kodu"),
    "urun": ("ürün adı", "ürün", "ürün ismi"),
    "kategori": ("kategori",),
    "adet": ("adet", "miktar"),
    "tutar": ("satış tutarı", "tutar", "faturalanacak tutar", "satır tutarı", "ödenen tutar"),
    "birim_fiyat": ("birim fiyat", "satış fiyatı"),
    "indirim": ("satıcı indirimi", "indirim", "satıcı indirim tutarı"),
    "komisyon_orani": ("komisyon oranı", "komisyon %"),
    "komisyon": ("komisyon", "komisyon tutarı"),
    "kargo": ("kargo", "kargo tutarı", "kargo bedeli"),
    "kdv": ("kdv oranı", "kdv"),
    "durum": ("durum", "sipariş durumu", "statü"),
}
MALIYET_ALANLARI = {
    "kod": ("stok kodu", "barkod", "ürün kodu", "sku", "model kodu"),
    "maliyet": ("birim maliyet", "maliyet", "alış fiyatı", "birim alış"),
    "kdv": ("kdv oranı", "kdv"),
    "desi": ("desi",),
    "ek": ("ek maliyet", "ambalaj", "paketleme"),
}


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Satir:
    siparis: str
    tarih: str
    kanal: str
    kod: str
    urun: str
    kategori: str
    adet: Decimal
    tutar: Decimal                  # KDV dahil, satıcı indirimi düşülmüş tahsil edilen tutar
    kdv_orani: Decimal
    durum: str
    komisyon_orani: Decimal | None
    komisyon: Decimal | None        # KDV dahil, dosyada verildiyse
    kargo: Decimal | None           # sipariş başına KDV dahil, dosyada verildiyse
    sonuc: dict = field(default_factory=dict)


def durum_turu(d: str) -> str:
    d = kucuk(d)
    if "iade" in d:
        return "İade"
    if "iptal" in d:
        return "İptal"
    return "Satış"


def kanallari_oku(yol: Path | None) -> dict:
    ham = json.loads(yol.read_text(encoding="utf-8")) if yol else {}
    kanallar = {}
    for ad, ayar in ham.items():
        k = dict(KANAL_VARSAYILAN)
        k.update(ayar)
        kanallar[ad] = k
    return kanallar


def maliyetleri_oku(yol: Path | None) -> dict:
    if not yol:
        return {}
    s = tablo_oku(yol)
    k = sutunlar(s[0], MALIYET_ALANLARI)
    if k["kod"] is None or k["maliyet"] is None:
        raise SystemExit(f"Maliyet listesinde Stok Kodu/Barkod ve Birim Maliyet sütunları gerekli. Başlıklar: {s[0]}")
    al = lambda r, a: r[k[a]] if k[a] is not None and k[a] < len(r) else None  # noqa: E731
    return {str(al(r, "kod")).strip(): {"maliyet": para(al(r, "maliyet")), "kdv": oran(al(r, "kdv")) if al(r, "kdv") not in (None, "") else None,
                                         "desi": para(al(r, "desi")), "ek": para(al(r, "ek"))}
            for r in s[1:] if al(r, "kod") not in (None, "")}


def siparisleri_oku(yol: Path, varsayilan_kanal: str) -> list[Satir]:
    s = tablo_oku(yol)
    k = sutunlar(s[0], SIPARIS_ALANLARI)
    if k["siparis"] is None or k["kod"] is None or (k["tutar"] is None and k["birim_fiyat"] is None):
        raise SystemExit(f"Sipariş dosyasında Sipariş No, Stok Kodu/Barkod ve Satış Tutarı sütunları gerekli. Başlıklar: {s[0]}")
    al = lambda r, a: r[k[a]] if k[a] is not None and k[a] < len(r) else None  # noqa: E731
    satirlar = []
    for r in s[1:]:
        if al(r, "siparis") in (None, ""):
            continue
        adet = para(al(r, "adet")) or Decimal(1)
        tutar = para(al(r, "tutar")) if k["tutar"] is not None else para(al(r, "birim_fiyat")) * adet
        tutar -= para(al(r, "indirim"))
        satirlar.append(Satir(
            str(al(r, "siparis")).strip(), str(al(r, "tarih") or ""), str(al(r, "kanal") or varsayilan_kanal).strip(),
            str(al(r, "kod")).strip(), str(al(r, "urun") or "").strip(), str(al(r, "kategori") or "").strip(), adet, tutar,
            oran(al(r, "kdv")) if al(r, "kdv") not in (None, "") else None, durum_turu(str(al(r, "durum") or "")),
            oran(al(r, "komisyon_orani")) if al(r, "komisyon_orani") not in (None, "") else None,
            para(al(r, "komisyon")) if al(r, "komisyon") not in (None, "") else None,
            para(al(r, "kargo")) if al(r, "kargo") not in (None, "") else None))
    return satirlar


# ----------------------------------------------------------------------------
# Hesap
# ----------------------------------------------------------------------------

def kargo_ucreti(kanal: dict, desi: Decimal) -> Decimal:
    tarife = sorted((para(ust), para(ucret)) for ust, ucret in kanal["desi_tarifesi"].items())
    if tarife:                      # desi bilinmiyorsa (0) en düşük kademe uygulanır
        for ust, ucret in tarife:
            if desi <= ust:
                return ucret
        return tarife[-1][1]
    return para(kanal["kargo_bedeli"])


def satir_hesapla(x: Satir, kanal: dict, maliyet: dict | None, kargo_payi: Decimal, hizmet_payi: Decimal, iade_kargo_payi: Decimal) -> dict:
    h_kdv = oran(kanal["hizmet_kdv_orani"])
    kdv = x.kdv_orani if x.kdv_orani is not None else (maliyet or {}).get("kdv") or Decimal("0.20")
    net_satis = x.tutar / (1 + kdv)
    satis_kdv = x.tutar - net_satis

    k_oran = x.komisyon_orani if x.komisyon_orani is not None else \
        oran(kanal["kategoriler"].get(x.kategori, kanal["komisyon_orani"]))
    if x.komisyon is not None:                       # dosyadaki gerçek komisyon (KDV dahil kabul edilir)
        kom_dahil = x.komisyon
    elif kanal["komisyon_kdv_dahil"]:
        kom_dahil = x.tutar * k_oran
    else:
        kom_dahil = net_satis * k_oran * (1 + h_kdv)
    kom_net = kom_dahil / (1 + h_kdv)

    birim_maliyet = (maliyet or {}).get("maliyet", SIFIR) + (maliyet or {}).get("ek", SIFIR)
    urun_maliyeti = birim_maliyet * x.adet

    if x.durum == "İptal":
        return {"net_satis": SIFIR, "satis_kdv": SIFIR, "komisyon": SIFIR, "komisyon_kdv": SIFIR, "kargo": SIFIR, "hizmet": SIFIR,
                "urun_maliyeti": SIFIR, "kar": SIFIR, "marj": None, "odeme": SIFIR, "stopaj": SIFIR, "kdv_etkisi": SIFIR,
                "k_oran": k_oran, "maliyet_yok": maliyet is None}

    kargo_dahil = kargo_payi
    hizmet_dahil = hizmet_payi
    if x.durum == "İade":
        # Satış, komisyon ve ürün maliyeti geri döner (ürün stoğa girer); gidiş + dönüş kargosu ve (ayara göre)
        # hizmet bedeli kayıptır.
        kargo_dahil = kargo_payi + iade_kargo_payi
        hizmet_dahil = hizmet_payi if kanal["iadede_hizmet_bedeli_kalir"] else SIFIR
        net_satis = satis_kdv = kom_dahil = kom_net = urun_maliyeti = SIFIR
    kargo_net = kargo_dahil / (1 + h_kdv)
    hizmet_net = hizmet_dahil / (1 + h_kdv)
    kar = net_satis - kom_net - kargo_net - hizmet_net - urun_maliyeti
    stopaj = net_satis * oran(kanal["stopaj_orani"])
    odeme = x.tutar - kom_dahil - kargo_dahil - hizmet_dahil - stopaj if x.durum == "Satış" else -(kargo_dahil + hizmet_dahil)
    # KDV etkisi: satış KDV'si − indirilecek KDV (komisyon, kargo, hizmet ve ürün alış KDV'si)
    alis_kdv = urun_maliyeti * kdv
    kdv_etkisi = satis_kdv - (kom_dahil - kom_net) - (kargo_dahil - kargo_net) - (hizmet_dahil - hizmet_net) - alis_kdv
    return {"net_satis": net_satis, "satis_kdv": satis_kdv, "komisyon": kom_net, "komisyon_kdv": kom_dahil - kom_net,
            "kargo": kargo_net, "hizmet": hizmet_net, "urun_maliyeti": urun_maliyeti, "kar": kar,
            "marj": None if not net_satis else kar / net_satis, "odeme": odeme, "stopaj": stopaj, "kdv_etkisi": kdv_etkisi,
            "k_oran": k_oran, "maliyet_yok": maliyet is None}


def basabas_fiyat(kanal: dict, kdv: Decimal, k_oran: Decimal, birim_maliyet: Decimal, kargo: Decimal, hizmet: Decimal) -> Decimal | None:
    """Tek adetlik siparişte kârı sıfırlayan KDV dahil satış fiyatı."""
    h_kdv = oran(kanal["hizmet_kdv_orani"])
    sabit = (kargo + hizmet) / (1 + h_kdv) + birim_maliyet
    if kanal["komisyon_kdv_dahil"]:
        katsayi = 1 / (1 + kdv) - k_oran / (1 + h_kdv)
    else:
        katsayi = (1 - k_oran) / (1 + kdv)
    return None if katsayi <= 0 else sabit / katsayi


def calistir(siparis_yolu: Path, maliyet_yolu: Path | None, kanal_yolu: Path | None, cikti: Path, varsayilan_kanal: str = "Pazaryeri") -> dict:
    satirlar = siparisleri_oku(siparis_yolu, varsayilan_kanal)
    maliyetler = maliyetleri_oku(maliyet_yolu)
    kanallar = kanallari_oku(kanal_yolu)
    uyarilar = []
    for ad in sorted({x.kanal for x in satirlar} - set(kanallar)):
        uyarilar.append(f"'{ad}' kanalı için ayar yok: komisyon, kargo ve hizmet bedeli 0 kabul edildi (kanallar.json'a ekleyin)")
        kanallar[ad] = dict(KANAL_VARSAYILAN)

    # Sipariş başına giderleri (kargo, hizmet bedeli) satırlara satış tutarı oranında dağıt
    siparisler: "OrderedDict[tuple, list[Satir]]" = OrderedDict()
    for x in satirlar:
        siparisler.setdefault((x.kanal, x.siparis), []).append(x)
    for (kanal_adi, _), liste in siparisler.items():
        kanal = kanallar[kanal_adi]
        dosya_kargo = [x.kargo for x in liste if x.kargo is not None]
        if dosya_kargo:
            kargo = max(dosya_kargo)       # sipariş dökümünde kargo her satırda tekrarlanabilir
        else:
            desi = sum((maliyetler.get(x.kod, {}).get("desi", SIFIR) * x.adet for x in liste), SIFIR)
            kargo = kargo_ucreti(kanal, desi)
        hizmet = para(kanal["hizmet_bedeli"])
        iade_kargo = para(kanal["iade_kargo_bedeli"]) if kanal["iade_kargo_bedeli"] not in (None, "") else kargo
        toplam = sum((x.tutar for x in liste), SIFIR) or Decimal(len(liste))
        for x in liste:
            pay = (x.tutar or Decimal(1)) / toplam
            m = maliyetler.get(x.kod)
            x.sonuc = satir_hesapla(x, kanal, m, kargo * pay, hizmet * pay, iade_kargo * pay)
            x.sonuc["siparis_kargo"], x.sonuc["siparis_hizmet"] = kargo, hizmet
    eksik = sorted({x.kod for x in satirlar if x.sonuc["maliyet_yok"] and x.durum != "İptal"})
    if eksik:
        uyarilar.append(f"Maliyeti bilinmeyen {len(eksik)} ürün (kâr olduğundan yüksek görünür): {', '.join(eksik[:10])}")

    # Ürün özeti
    urunler: dict[tuple, dict] = {}
    for x in satirlar:
        if x.durum == "İptal":
            continue
        u = urunler.setdefault((x.kanal, x.kod), defaultdict(Decimal) | {"urun": x.urun, "satir": x})
        u["adet"] += x.adet if x.durum == "Satış" else SIFIR
        u["iade"] += x.adet if x.durum == "İade" else SIFIR
        for k in ("net_satis", "komisyon", "kargo", "hizmet", "urun_maliyeti", "kar", "odeme"):
            u[k] += x.sonuc[k]
    for (kanal_adi, kod), u in urunler.items():
        x, m = u["satir"], maliyetler.get(kod)
        kdv = x.kdv_orani if x.kdv_orani is not None else (m or {}).get("kdv") or Decimal("0.20")
        kanal = kanallar[kanal_adi]
        tek_kargo = kargo_ucreti(kanal, (m or {}).get("desi", SIFIR)) if x.kargo is None else x.kargo
        u["basabas"] = None if m is None else basabas_fiyat(kanal, kdv, x.sonuc["k_oran"], m["maliyet"] + m["ek"], tek_kargo,
                                                             para(kanal["hizmet_bedeli"]))
        u["ort_fiyat"] = None if not u["adet"] else sum((y.tutar for y in satirlar if y.kanal == kanal_adi and y.kod == kod
                                                         and y.durum == "Satış"), SIFIR) / u["adet"]
        u["marj"] = None if not u["net_satis"] else u["kar"] / u["net_satis"]
        u["zarar"] = u["kar"] < 0
    _rapor(satirlar, urunler, cikti, uyarilar)
    toplam = defaultdict(Decimal)
    for x in satirlar:
        for k in ("net_satis", "komisyon", "kargo", "hizmet", "urun_maliyeti", "kar", "odeme", "stopaj", "kdv_etkisi"):
            toplam[k] += x.sonuc[k]
    return {"satirlar": satirlar, "urunler": urunler, "toplam": toplam, "uyarilar": uyarilar}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
PARA = "#,##0.00"


def tl(x) -> str:
    return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = f"C{ws.max_row + 1}"


def _sayilar(ws, para_sutunlari, yuzde_sutunlari=()):
    for c in para_sutunlari:
        ws.cell(ws.max_row, c).number_format = PARA
    for c in yuzde_sutunlari:
        ws.cell(ws.max_row, c).number_format = "0.0%"


def f(x):
    return None if x is None else float(yuvarla(x) if isinstance(x, Decimal) else x)


def _rapor(satirlar, urunler, cikti, uyarilar):
    wb = Workbook()
    u = wb.active
    u.title = "Ürün Kârlılığı"
    _baslik(u, ["Kanal", "Stok Kodu", "Ürün", "Satılan Adet", "İade Adet", "Ort. Satış Fiyatı (KDV dahil)", "Net Satış (KDV hariç)",
                "Komisyon", "Kargo", "Hizmet Bedeli", "Ürün Maliyeti", "Net Kâr", "Marj", "Başabaş Fiyat (KDV dahil)", "Durum"],
            (14, 14, 34, 11, 10, 16, 16, 14, 12, 12, 14, 14, 9, 16, 30))
    for (kanal, kod), x in sorted(urunler.items(), key=lambda i: i[1]["kar"]):
        durum = "ZARARINA SATIŞ" if x["zarar"] else ""
        if x["basabas"] is not None and x["ort_fiyat"] and x["ort_fiyat"] < x["basabas"]:
            durum = (durum + " · " if durum else "") + "fiyat başabaşın altında"
        u.append([kanal, kod, x["urun"], f(x["adet"]), f(x["iade"]), f(x["ort_fiyat"]), f(x["net_satis"]), f(x["komisyon"]), f(x["kargo"]),
                  f(x["hizmet"]), f(x["urun_maliyeti"]), f(x["kar"]), f(x["marj"]), f(x["basabas"]), durum])
        _sayilar(u, (6, 7, 8, 9, 10, 11, 12, 14), (13,))
        if durum:
            u.cell(u.max_row, 12).fill = u.cell(u.max_row, 15).fill = KIRMIZI
    u.auto_filter.ref = u.dimensions

    d = wb.create_sheet("Sipariş Detayı")
    _baslik(d, ["Kanal", "Sipariş No", "Tarih", "Stok Kodu", "Ürün", "Adet", "Durum", "Satış Tutarı (KDV dahil)", "Satış KDV",
                "Net Satış", "Komisyon %", "Komisyon (KDV hariç)", "Kargo Payı (KDV hariç)", "Hizmet Payı (KDV hariç)",
                "Ürün Maliyeti", "Net Kâr", "Marj", "Pazaryeri Ödemesi", "Stopaj (%1)", "KDV Etkisi"],
            (14, 16, 12, 14, 30, 7, 9, 15, 12, 13, 10, 14, 14, 14, 13, 13, 9, 15, 12, 12))
    for x in satirlar:
        s = x.sonuc
        d.append([x.kanal, x.siparis, x.tarih, x.kod, x.urun, f(x.adet), x.durum, f(x.tutar), f(s["satis_kdv"]), f(s["net_satis"]),
                  f(s["k_oran"]), f(s["komisyon"]), f(s["kargo"]), f(s["hizmet"]), f(s["urun_maliyeti"]), f(s["kar"]), f(s["marj"]),
                  f(s["odeme"]), f(s["stopaj"]), f(s["kdv_etkisi"])])
        _sayilar(d, (8, 9, 10, 12, 13, 14, 15, 16, 18, 19, 20), (11, 17))
        if s["kar"] < 0:
            d.cell(d.max_row, 16).fill = KIRMIZI
    d.auto_filter.ref = d.dimensions

    k = wb.create_sheet("Kanal Özeti", 0)
    k.append(["Pazaryeri Sipariş Kârlılığı — Kanal Özeti"])
    k["A1"].font = Font(bold=True, size=13)
    k.append([])
    k.append(["Kanal", "Sipariş", "Satış Adedi", "İade Adedi", "Net Satış", "Komisyon", "Kargo", "Hizmet Bedeli", "Ürün Maliyeti",
              "Net Kâr", "Marj", "Pazaryeri Ödemesi", "Stopaj", "KDV Etkisi"])
    for h in k[3]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    kanallar = OrderedDict()
    for x in satirlar:
        t = kanallar.setdefault(x.kanal, defaultdict(Decimal) | {"siparis": set()})
        if x.durum != "İptal":
            t["siparis"].add(x.siparis)
        t["adet"] += x.adet if x.durum == "Satış" else SIFIR
        t["iade"] += x.adet if x.durum == "İade" else SIFIR
        for a in ("net_satis", "komisyon", "kargo", "hizmet", "urun_maliyeti", "kar", "odeme", "stopaj", "kdv_etkisi"):
            t[a] += x.sonuc[a]
    for ad, t in kanallar.items():
        k.append([ad, len(t["siparis"]), f(t["adet"]), f(t["iade"]), f(t["net_satis"]), f(t["komisyon"]), f(t["kargo"]), f(t["hizmet"]),
                  f(t["urun_maliyeti"]), f(t["kar"]), f(t["kar"] / t["net_satis"]) if t["net_satis"] else None, f(t["odeme"]),
                  f(t["stopaj"]), f(t["kdv_etkisi"])])
        _sayilar(k, (5, 6, 7, 8, 9, 10, 12, 13, 14), (11,))
    zararli = [x for x in urunler.values() if x["zarar"]]
    k.append([])
    k.append([f"Zararına satılan ürün sayısı: {len(zararli)}"])
    for uy in uyarilar:
        k.append(["Uyarı", uy])
    for j, w in enumerate((22, 10, 11, 10, 15, 14, 13, 13, 15, 14, 9, 17, 12, 12), 1):
        k.column_dimensions[get_column_letter(j)].width = w

    b = wb.create_sheet("Bilgi")
    for s in [["Net kâr", "Net satış (KDV hariç) − komisyon − kargo − hizmet bedeli (KDV hariç) − ürün maliyeti (KDV hariç). Vergi öncesidir"],
              ["Komisyon", "komisyon_kdv_dahil=true: KDV dahil satış tutarı × oran, bu tutar KDV'yi içerir (Trendyol Satıcı Bilgi Merkezi). "
                           "false: KDV hariç tutar × oran + KDV"],
              ["Sipariş giderleri", "Kargo ve hizmet bedeli siparişteki satırlara satış tutarı oranında dağıtılır"],
              ["İade", "Satış, komisyon ve ürün maliyeti geri alınır (ürün stoğa döner); gidiş + dönüş kargosu ve hizmet bedeli kayıptır"],
              ["Stopaj", "E-ticaret stopajı %1 × KDV hariç satış (9284 sayılı CBK, 01.01.2025). Vergiden mahsup edilir; kârı etkilemez"],
              ["KDV etkisi", "Satış KDV'si − (komisyon, kargo, hizmet ve ürün alış KDV'si): bu satıştan doğan ödenecek KDV"],
              ["Başabaş fiyat", "Tek adetlik siparişte kârı sıfırlayan KDV dahil satış fiyatı"],
              ["Uyarı", "Komisyon, kargo ve hizmet bedelleri sık değişir; kanallar.json'u satıcı panelinizdeki güncel değerlerle güncelleyin"]]:
        b.append(s)
    b.column_dimensions["A"].width = 18
    b.column_dimensions["B"].width = 130
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description="Pazaryeri siparişlerinden sipariş ve ürün bazında net kâr hesaplar.")
    ap.add_argument("--siparisler", type=Path, default=BURASI / "ornek_veri" / "siparisler.csv",
                    help="Sipariş dökümü (.xlsx/.csv): Sipariş No, Stok Kodu/Barkod, Satış Tutarı (KDV dahil), [Kanal, Adet, Durum, ...]")
    ap.add_argument("--maliyetler", type=Path, default=BURASI / "ornek_veri" / "urun_maliyetleri.csv",
                    help="Ürün maliyetleri (.xlsx/.csv): Stok Kodu, Birim Maliyet (KDV hariç), [KDV Oranı, Desi, Ek Maliyet]")
    ap.add_argument("--kanallar", type=Path, default=BURASI / "ornek_veri" / "kanallar.json", help="Kanal ayarları (JSON)")
    ap.add_argument("--kanal", default="Pazaryeri", help="Sipariş dosyasında Kanal sütunu yoksa kullanılacak kanal adı")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "pazaryeri_karlilik.xlsx")
    a = ap.parse_args(argv)
    s = calistir(a.siparisler, a.maliyetler, a.kanallar, a.cikti, a.kanal)
    t = s["toplam"]
    zararli = [u for u in s["urunler"].values() if u["zarar"]]
    print(f"[OK] {len(s['satirlar'])} sipariş satırı · net satış {tl(t['net_satis'])} TL · net kâr {tl(t['kar'])} TL"
          + (f" (%{t['kar'] / t['net_satis'] * 100:.1f})".replace(".", ",") if t["net_satis"] else ""))
    print(f"[OK] Pazaryeri ödemesi {tl(t['odeme'])} TL · stopaj {tl(t['stopaj'])} TL")
    for u in sorted(zararli, key=lambda u: u["kar"])[:10]:
        print(f"[!] Zararına: {u['satir'].kod} {u['urun'][:30]} · kâr {tl(u['kar'])} TL"
              + (f" · başabaş {tl(u['basabas'])} TL" if u["basabas"] is not None else ""))
    for uy in s["uyarilar"]:
        print(f"[!] {uy}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
