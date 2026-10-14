import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def o(no, kazanc, maliyet, uyg, bolum="B"):
    return {"no": no, "oneren": "", "bolum": bolum, "konu": "Konu", "kazanc": D(kazanc), "maliyet": D(maliyet), "uyg": D(uyg)}


def veri():
    return [o("k1", 120000, 20000, 4), o("k2", 8000, 12000, 5, "C"), o("k3", 50000, 0, 3), o("k4", 0, 5000, 2, "C"), o("k5", 250000, 600000, 2, "c")]


class PuanTesti(unittest.TestCase):
    def test_puanlar(self):
        r = main.puanla(veri())
        self.assertEqual([(x["sira"], x["no"], x["kp"], x["ay"], x["gp"], x["puan"], x["sinif"], x["net"]) for x in r["oneriler"]],
                         [(1, "k1", 4, D("2.0"), 5, D("86.00"), "Öncelikli", D(100000)), (2, "k3", 3, D(0), 5, D("72.00"), "Öncelikli", D(50000)), (3, "k5", 5, D("28.8"), 1, D("64.00"), "Değerlendir", D(-350000)),
                          (4, "k2", 1, D("18.0"), 2, D("42.00"), "Düşük öncelik", D(-4000)), (5, "k4", 1, None, 1, D("24.00"), "Düşük öncelik", D(-5000))])
        self.assertEqual(r["ozet"], {"oneri": 5, "oncelikli": 2, "degerlendir": 1, "dusuk": 2, "kazanc": D(428000), "maliyet": D(637000), "ort_puan": D("57.60"), "onc_kazanc": D(170000), "onc_maliyet": D(20000)})
        self.assertEqual([(b["bolum"], b["oneri"], b["oncelikli"], b["degerlendir"], b["dusuk"], b["kazanc"], b["maliyet"], b["ort_puan"]) for b in r["bolumler"]],
                         [("C", 3, 0, 1, 2, D(258000), D(617000), D("43.33")), ("B", 2, 2, 0, 0, D(170000), D(20000), D("79.00"))])
        self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]], [("Bilgi", "Parasal kazancı yazılmamış öneri", "k4")])

    def test_sinirlar_ve_secenekler(self):
        kp = lambda x: main.puanla([o("a", x, 0, 3)])["oneriler"][0]["kp"]   # noqa: E731
        gp = lambda m: main.puanla([o("a", 12000, m, 3)])["oneriler"][0]["gp"]   # noqa: E731
        self.assertEqual(([kp(x) for x in (9999, 10000, 49999, 50000, 100000, 249999, 250000)], [gp(m) for m in (3000, 3001, 6000, 6001, 12000, 12001, 24000, 24001)]), ([1, 2, 2, 3, 4, 4, 5], [5, 4, 4, 3, 3, 2, 2, 1]))
        s = lambda **kw: {x["no"]: (x["puan"], x["sinif"]) for x in main.puanla(veri(), **kw)["oneriler"]}   # noqa: E731
        self.assertEqual((s(agirlik=(D(0), D(0), D(1)))["k2"], s(agirlik=(D(1), D(0), D(0)))["k5"], s(oncelikli=D(86), degerlendir=D(72))["k1"], s(oncelikli=D("86.01"), degerlendir=D("72.01"))["k3"]),
                         ((D("100.00"), "Öncelikli"), (D("100.00"), "Öncelikli"), (D("86.00"), "Öncelikli"), (D("72.00"), "Düşük öncelik")))
        r = main.puanla([o("x", 300000, 400000, 5), o("b", 50000, 0, 3), o("a", 50000, 0, 3)], agirlik=(D(100), D(0), D(0)))
        self.assertEqual(([x["no"] for x in r["oneriler"]], [u["tur"] for u in r["uyarilar"]]), (["x", "a", "b"], ["Öncelikli ama ilk yıl kendini ödemiyor"]))
        with self.assertRaises(ValueError):
            main.puanla([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "kaizen_onerileri.csv", cikti)
            self.assertEqual([(x["no"], x["puan"], x["sinif"], x["ay"]) for x in r["oneriler"]],
                             [("K-01", D("86.00"), "Öncelikli", D("3.0")), ("K-03", D("76.00"), "Öncelikli", D("2.0")), ("K-02", D("70.00"), "Öncelikli", D("1.5")), ("K-05", D("70.00"), "Öncelikli", D("15.0")),
                              ("K-04", D("60.00"), "Değerlendir", D("2.0")), ("K-07", D("48.00"), "Düşük öncelik", D("16.0")), ("K-06", D("48.00"), "Düşük öncelik", D("46.7")), ("K-08", D("32.00"), "Düşük öncelik", None)])
            o_ = r["ozet"]
            self.assertEqual((o_["oneri"], o_["oncelikli"], o_["degerlendir"], o_["dusuk"], o_["kazanc"], o_["maliyet"], o_["onc_kazanc"], o_["onc_maliyet"]), (8, 4, 1, 3, D(698000), D(854500), D(584000), D(458000)))
            self.assertEqual([(b["bolum"], b["oneri"], b["kazanc"]) for b in r["bolumler"]], [("Pres", 2, D(500000)), ("Montaj", 2, D(114000)), ("Depo", 2, D(9000)), ("Boyahane", 1, D(60000)), ("Kalite", 1, D(15000))])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]], [("Orta", "Öncelikli ama ilk yıl kendini ödemiyor", "K-05"), ("Bilgi", "Parasal kazancı yazılmamış öneri", "K-08")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Öneriler", "Özet", "Bölümler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "k.csv"
            y.write_text("Kaizen No;Çalışan;Birim;Öneri;Yıllık Tasarruf;Maliyet;Kolaylık\nK1;Ali;Pres;Aparat;12.500,50;1.000;4\n;;;Düzen;5000;;\nK3;;;X;abc;;\nK4;;;X;10;;6\nK5;;;X;10;-1;\n;;;;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["oneren"], x["bolum"], x["konu"], x["kazanc"], x["maliyet"], x["uyg"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("K1", "Ali", "Pres", "Aparat", D("12500.50"), D(1000), D(4)), ("Satır 3", "", "(bölüm yok)", "Düzen", D(5000), D(0), D(3))],
                              [("Uygulanabilirlik puanı yok", "Satır 3"), ("Öneri okunamadı", "K3"), ("Öneri okunamadı", "K4"), ("Öneri okunamadı", "K5")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--oneriler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--kazanc", "25000;100000;250000;1000000", "--geri-odeme", "2;4;8;16", "--agirlik", "40;20;40", "--oncelikli", "75", "--degerlendir", "40"]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--kazanc", "1;2;3"], ["--kazanc", "4;3;2;1"], ["--geri-odeme", "a;b;c;d"], ["--agirlik", "0;0;0"], ["--oncelikli", "40"], ["--degerlendir", "10"], ["--oncelikli", "101"])], [2] * 7)
            self.assertEqual(main.main(["--oneriler", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Konu;Bölüm\nA;B\n", "Konu;Yıllık Kazanç\n;\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
