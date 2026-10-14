"""
Kaizen Önerisi Değerlendirme — Workers / Workless kod bloğu
Endüstri Mühendisliği ve Süreç Geliştirme › Endüstri Mühendisi

Çalışan iyileştirme önerilerini kazanç, geri ödeme süresi ve uygulanabilirliğe göre puanlar ve sıralar.
  Kazanç puanı (1–5)        = yıllık kazancın geçtiği eşik sayısı + 1          (--kazanc 10000;50000;100000;250000)
  Geri ödeme süresi (ay)    = uygulama maliyeti ÷ (yıllık kazanç ÷ 12);   maliyet sıfırsa 0, kazanç sıfırsa hesaplanmaz
  Geri ödeme puanı (1–5)    = 5 − sürenin aştığı eşik sayısı                  (--geri-odeme 3;6;12;24 ay);  hesaplanamıyorsa 1
  Uygulanabilirlik (1–5)    = dosyada sizin verdiğiniz puan
  Toplam puan (20–100)      = ağırlıklı ortalama × 20                          (--agirlik 50;30;20 → kazanç; geri ödeme; uygulanabilirlik)
  Sınıf                     = Öncelikli (puan ≥ --oncelikli, 70) · Değerlendir (≥ --degerlendir, 50) · Düşük öncelik
  Net ilk yıl kazancı       = yıllık kazanç − uygulama maliyeti
Kazanç ve maliyet tahminini siz yaparsınız; kod parasal olmayan faydaları (güvenlik, ergonomi, kalite) puanlamaz.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"no": ("oneri no", "kaizen no", "no"), "oneren": ("oneren", "oneri sahibi", "calisan"), "bolum": ("bolum", "birim", "departman"), "konu": ("konu", "oneri", "aciklama"),
         "kazanc": ("yillik kazanc", "kazanc", "tahmini yillik kazanc", "yillik tasarruf"), "maliyet": ("uygulama maliyeti", "maliyet", "yatirim tutari"), "uyg": ("uygulanabilirlik", "uygulanabilirlik puani", "kolaylik")}
SINIF_RENK = {"Öncelikli": YESIL, "Değerlendir": SARI, "Düşük öncelik": None}
SIFIR = Decimal(0)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("konu", "kazanc")):
        konu, no = metin(r["konu"]), metin(r["no"]) or f"Satır {r['_satir']}"
        if not konu and not metin(r["kazanc"]):
            continue
        kazanc, maliyet, uyg = para(r["kazanc"]) if metin(r["kazanc"]) else SIFIR, para(r["maliyet"]) if metin(r["maliyet"]) else SIFIR, para(r["uyg"]) if metin(r["uyg"]) else None
        if kazanc is None or kazanc < 0 or maliyet is None or maliyet < 0 or (metin(r["uyg"]) and (uyg is None or not 1 <= uyg <= 5)):
            uy.append(uyari("Yüksek", "Öneri okunamadı", no, f"Satır {r['_satir']}: kazanç, maliyet ya da uygulanabilirlik (1–5) okunamadı. Satır alınmadı."))
            continue
        if uyg is None:
            uy.append(uyari("Orta", "Uygulanabilirlik puanı yok", no, "Puan yazılmamış; orta değer (3) kullanıldı."))
        l.append({"no": no, "oneren": metin(r["oneren"]), "bolum": metin(r["bolum"]) or "(bölüm yok)", "konu": konu, "kazanc": kazanc, "maliyet": maliyet, "uyg": uyg if uyg is not None else Decimal(3)})
    return l, uy


def puanla(oneriler: list[dict], kazanc_esik: tuple[Decimal, ...] = (Decimal(10000), Decimal(50000), Decimal(100000), Decimal(250000)), geri_esik: tuple[Decimal, ...] = (Decimal(3), Decimal(6), Decimal(12), Decimal(24)),
           agirlik: tuple[Decimal, ...] = (Decimal(50), Decimal(30), Decimal(20)), oncelikli: Decimal = Decimal(70), degerlendir: Decimal = Decimal(50)) -> dict:
    if not oneriler:
        raise ValueError("Değerlendirilecek öneri yok.")
    uy = []
    for o in oneriler:
        o["kp"] = sum(1 for e in kazanc_esik if o["kazanc"] >= e) + 1
        o["ay"] = SIFIR if o["maliyet"] == 0 else yuvarla(o["maliyet"] * 12 / o["kazanc"], "0.1") if o["kazanc"] > 0 else None
        o["gp"] = 1 if o["ay"] is None else 5 - sum(1 for e in geri_esik if o["maliyet"] * 12 > e * o["kazanc"])
        o["puan"] = yuvarla((o["kp"] * agirlik[0] + o["gp"] * agirlik[1] + o["uyg"] * agirlik[2]) / sum(agirlik) * 20)
        o["sinif"] = "Öncelikli" if o["puan"] >= oncelikli else "Değerlendir" if o["puan"] >= degerlendir else "Düşük öncelik"
        o["net"] = o["kazanc"] - o["maliyet"]
        if o["kazanc"] == 0:
            uy.append(uyari("Bilgi", "Parasal kazancı yazılmamış öneri", o["no"], f"{o['konu']}: kazanç sıfır olduğu için puanı düşük çıktı. Güvenlik, ergonomi ya da kalite faydası varsa ayrıca değerlendirin."))
        elif o["net"] < 0 and o["sinif"] == "Öncelikli":
            uy.append(uyari("Orta", "Öncelikli ama ilk yıl kendini ödemiyor", o["no"], f"{o['konu']}: yıllık kazanç {o['kazanc']:,.2f}, maliyet {o['maliyet']:,.2f}; geri ödeme {o['ay']} ay. Bütçe onayı gerektirebilir."))
    sirali = sorted(oneriler, key=lambda o: (-o["puan"], -o["net"], katla(o["no"])))
    for n, o in enumerate(sirali, 1):
        o["sira"] = n

    def ozet(l: list[dict]) -> dict:
        return {"oneri": len(l), "oncelikli": sum(1 for o in l if o["sinif"] == "Öncelikli"), "degerlendir": sum(1 for o in l if o["sinif"] == "Değerlendir"), "dusuk": sum(1 for o in l if o["sinif"] == "Düşük öncelik"),
                "kazanc": sum((o["kazanc"] for o in l), SIFIR), "maliyet": sum((o["maliyet"] for o in l), SIFIR), "ort_puan": yuvarla(sum((o["puan"] for o in l), SIFIR) / len(l))}

    grup = {}
    for o in oneriler:
        grup.setdefault(katla(o["bolum"]), []).append(o)
    onc = [o for o in oneriler if o["sinif"] == "Öncelikli"]
    return {"oneriler": sirali, "ozet": {**ozet(oneriler), "onc_kazanc": sum((o["kazanc"] for o in onc), SIFIR), "onc_maliyet": sum((o["maliyet"] for o in onc), SIFIR)},
            "bolumler": sorted(({"bolum": g[0]["bolum"], **ozet(g)} for g in grup.values()), key=lambda b: (-b["oneri"], -b["kazanc"], katla(b["bolum"]))), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, agirlik: tuple[Decimal, ...], oncelikli: Decimal, degerlendir: Decimal) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Öneriler"
    _baslik(ws, ["Sıra", "Sınıf", "Toplam Puan", "Öneri No", "Konu", "Bölüm", "Öneren", "Yıllık Kazanç", "Uygulama Maliyeti", "Net İlk Yıl", "Geri Ödeme (ay)", "Kazanç Puanı", "Geri Ödeme Puanı", "Uygulanabilirlik", "Karar (siz doldurun)"],
            (6, 15, 9, 10, 44, 16, 18, 15, 15, 15, 10, 9, 10, 12, 28))
    for o in r["oneriler"]:
        _satir(ws, [o["sira"], o["sinif"], o["puan"], o["no"], o["konu"], o["bolum"], o["oneren"], o["kazanc"], o["maliyet"], o["net"], o["ay"], o["kp"], o["gp"], o["uyg"], None], para_s=(3, 8, 9, 10),
               dolgu={2: SINIF_RENK[o["sinif"]], 10: KIRMIZI if o["net"] < 0 else None}, sar=(5,))
    ws.append([])
    ws.append([f"Ağırlıklar: kazanç {agirlik[0]}, geri ödeme {agirlik[1]}, uygulanabilirlik {agirlik[2]} · Öncelikli ≥ {oncelikli}, Değerlendir ≥ {degerlendir}"])
    ws = wb.create_sheet("Özet")
    o = r["ozet"]
    _baslik(ws, ["Gösterge", "Değer"], (44, 18))
    for etiket, deger, para_mi in (("Öneri", o["oneri"], False), ("Öncelikli", o["oncelikli"], False), ("Değerlendir", o["degerlendir"], False), ("Düşük öncelik", o["dusuk"], False), ("Ortalama puan", o["ort_puan"], True),
                                   ("Bütün önerilerin yıllık kazancı", o["kazanc"], True), ("Bütün önerilerin uygulama maliyeti", o["maliyet"], True), ("Öncelikli önerilerin yıllık kazancı", o["onc_kazanc"], True),
                                   ("Öncelikli önerilerin uygulama maliyeti", o["onc_maliyet"], True)):
        _satir(ws, [etiket, deger], para_s=(2,) if para_mi else ())
    ws = wb.create_sheet("Bölümler")
    _baslik(ws, ["Bölüm", "Öneri", "Öncelikli", "Değerlendir", "Düşük Öncelik", "Yıllık Kazanç", "Uygulama Maliyeti", "Ortalama Puan"], (24, 8, 10, 11, 12, 16, 16, 12))
    for b in r["bolumler"]:
        _satir(ws, [b["bolum"], b["oneri"], b["oncelikli"] or None, b["degerlendir"] or None, b["dusuk"] or None, b["kazanc"], b["maliyet"], b["ort_puan"]], para_s=(6, 7, 8))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, **kw) -> dict:
    oneriler, uy = oku(yol)
    r = puanla(oneriler, **kw)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    varsayilan = puanla.__defaults__
    rapor_yaz(cikti, r, kw.get("agirlik", varsayilan[2]), kw.get("oncelikli", varsayilan[3]), kw.get("degerlendir", varsayilan[4]))
    return r


def _liste(metni: str, adet: int) -> tuple[Decimal, ...] | None:
    d = tuple(para(x) for x in metni.split(";"))
    return d if len(d) == adet and all(x is not None and x >= 0 for x in d) else None


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Çalışan iyileştirme önerilerini kazanç, maliyet ve uygulanabilirliğe göre puanlar.")
    p.add_argument("--oneriler", type=Path, help="Öneri No, Öneren, Bölüm, Konu, Yıllık Kazanç, Uygulama Maliyeti, Uygulanabilirlik (1–5)")
    p.add_argument("--kazanc", default="10000;50000;100000;250000", help="Kazanç puanı eşikleri, artan dört tutar (varsayılan 10000;50000;100000;250000)")
    p.add_argument("--geri-odeme", default="3;6;12;24", help="Geri ödeme puanı eşikleri, artan dört süre, ay (varsayılan 3;6;12;24)")
    p.add_argument("--agirlik", default="50;30;20", help="Kazanç;geri ödeme;uygulanabilirlik ağırlıkları (varsayılan 50;30;20)")
    p.add_argument("--oncelikli", default="70", help="Öncelikli sınıfı için en düşük puan (varsayılan 70)")
    p.add_argument("--degerlendir", default="50", help="Değerlendir sınıfı için en düşük puan (varsayılan 50)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kaizen_degerlendirme.xlsx")
    a = p.parse_args(argv)
    yol = a.oneriler or ORNEK / "kaizen_onerileri.csv"
    ke, ge, ag, onc, deg = _liste(a.kazanc, 4), _liste(a.geri_odeme, 4), _liste(a.agirlik, 3), para(a.oncelikli), para(a.degerlendir)
    if ke is None or ge is None or ag is None or list(ke) != sorted(set(ke)) or list(ge) != sorted(set(ge)) or sum(ag) <= 0 or onc is None or deg is None or not 20 <= deg <= onc <= 100:
        print("[X] --kazanc ve --geri-odeme artan dört sayı; --agirlik üç sayı (toplamı sıfırdan büyük); 20 ≤ --degerlendir ≤ --oncelikli ≤ 100 olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, kazanc_esik=ke, geri_esik=ge, agirlik=ag, oncelikli=onc, degerlendir=deg)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {o['oneri']} öneri · öncelikli {o['oncelikli']}, değerlendir {o['degerlendir']}, düşük öncelik {o['dusuk']} · öncelikli önerilerin yıllık kazancı {o['onc_kazanc']:,.2f}, maliyeti {o['onc_maliyet']:,.2f}")
    for x in r["oneriler"][:3]:
        print(f"     {x['sira']}. {x['no']} · {x['konu']} · puan {x['puan']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
