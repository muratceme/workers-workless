import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
Q = D("0.0001")


def q(x):
    return x.quantize(Q)


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "c.xlsx"
        cls.s = main.calistir(ORNEK / "modeller.csv", ORNEK / "maliyet_kalemleri.csv", cls.cikti, dict(main.ORNEK_KUR))
        cls.r = {x.model.kod: x for x in cls.s["sonuclar"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_ts01_elle(self):
        x = self.r["TS-01"]
        kumas = D("0.22") * D("1.08") * D("5.10") + D("0.015") * D("1.10") * D("6.20")
        cm = D(12) * D("4.10") / D("0.65") / D("41.20")
        sabit = (D(450) + D(3000) / D("41.20")) / 12000
        self.assertEqual(q(x.gruplar["Kumaş"]), q(kumas))
        self.assertEqual(q(x.gruplar["CM"]), q(cm))
        self.assertEqual(q(x.gruplar["Sabit"]), q(sabit))
        self.assertEqual(q(x.toplam), q(x.uretim * D("1.08")))
        self.assertEqual(q(x.fob), q(x.toplam / D("0.85")))
        self.assertEqual(q(x.hedef_marj), q((D("4.20") * D("0.97") - x.toplam) / D("4.20")))
        self.assertEqual(q(x.hedef_dusus), q(x.toplam - D("4.20") * D("0.85")))

    def test_sonuclar(self):
        self.assertEqual({k: main.fmt(x.fob) for k, x in self.r.items()}, {"TS-01": "4,56", "SW-02": "8,92", "DR-03": "15,82"})
        self.assertIsNone(self.r["SW-02"].hedef_dusus or None)

    def test_duyarlilik(self):
        d = {ad: fob for ad, _, fob in self.r["TS-01"].duyarlilik}
        self.assertGreater(d["USD/TL kuru %-5"], self.r["TS-01"].fob)
        self.assertLess(d["USD/TL kuru %+5"], self.r["TS-01"].fob)
        self.assertEqual(main.fmt(d["Kumaş fiyatı %+10"]), "4,72")

    def test_uyarilar(self):
        u = {(x["tur"], x["model"]): x for x in self.s["uyarilar"]}
        self.assertIn(("Kur eksik", "DR-03"), u)
        self.assertIn(("Fiyat eksik", "SW-02"), u)
        self.assertEqual(u[("Hedef fiyat altında", "DR-03")]["onem"], "Yüksek")
        self.assertEqual(u[("Hedef fiyat altında", "TS-01")]["onem"], "Orta")
        self.assertNotIn(("Hedef fiyat altında", "SW-02"), u)
        askı = next(k for k in self.r["DR-03"].kalemler if k.ad == "Askı")
        self.assertIsNone(askı.tutar)

    def test_excel(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Özet", "Föy TS-01", "Föy SW-02", "Föy DR-03", "Duyarlılık", "Uyarılar"])
        self.assertEqual(wb["Özet"]["Q2"].value, 4.56)
        self.assertEqual(wb["Özet"]["U1"].value, "Karar / Not")


class KuralTesti(unittest.TestCase):
    def test_yardimcilar(self):
        self.assertEqual(main.yuzde("8"), D("0.08"))
        self.assertEqual(main.yuzde("0,08"), D("0.08"))
        self.assertEqual(main.sayi("1.250,50"), D("1250.50"))
        self.assertEqual(main.kur_coz(["USD=41,20"]), {"USD": D("41.20")})
        self.assertEqual(main.grup_bul("Baskı"), "Fason İşlem")
        self.assertEqual(main.cevir(D(100), "EUR", "USD", {"EUR": D(48), "USD": D(40)}), D(120))
        with self.assertRaises(ValueError):
            main.kur_coz(["USD:41"])

    def test_tl_teklif_ve_eksikler(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t / "m.csv").write_text("Model;Sipariş Adedi;Teklif Para Birimi;Kâr Marjı %\nA;1000;TL;20\nB;100;TL;10\n", encoding="utf-8")
            (t / "k.csv").write_text("Model;Grup;Kalem;Miktar;Fire %;Birim Fiyat;Para Birimi\nA;Kumaş;Penye;1;0;100;TL\nZ;Aksesuar;Etiket;1;;1;TL\n",
                                     encoding="utf-8")
            s = main.calistir(t / "m.csv", t / "k.csv", t / "o.xlsx", {})
            self.assertEqual(s["sonuclar"][0].fob, D(125))
            turler = {(u["tur"], u["model"]) for u in s["uyarilar"]}
            for b in (("Kalem yok", "B"), ("Model tanımsız", "Z"), ("CM kalemi yok", "A"), ("Kumaş firesi 0", "A")):
                self.assertIn(b, turler)

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--kalemler", str(Path(tmp) / "yok.csv")]), 1)
            self.assertEqual(main.main(["--kur", "USD41", "--cikti", str(Path(tmp) / "x.xlsx")]), 1)


if __name__ == "__main__":
    unittest.main()
