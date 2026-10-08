"""
Ürün Maliyet Hesaplama (Costing) — Workers / Workless kod bloğu
Tekstil ve Konfeksiyon › Müşteri Temsilciliği (Merchandising) › Development Merchandiser

Hazır giyim maliyet föyü (cost sheet) hazırlar ve FOB teklif fiyatını hesaplar:
  - Kumaş: tüketim × (1 + fire %) × birim fiyat (kg veya metre)
  - Aksesuar ve fason işlemler (baskı, nakış, yıkama): miktar × (1 + fire %) × birim fiyat
  - CM (kesim-dikim-ütü paket): SAM (dk) × dakika maliyeti ÷ hat verimliliği
  - Diğer adet başı giderler (nakliye vb.) ve sipariş başı sabit giderler (test, numune) ÷ sipariş adedi
  - Genel gider = üretim maliyeti × GG %
  - FOB teklif fiyatı = toplam maliyet ÷ (1 − kâr marjı % − komisyon %); marj ve komisyon satış fiyatı üzerinden
Farklı para birimindeki kalemler verilen kurlarla teklif para birimine çevrilir. Müşteri hedef fiyatı verilirse
hedef fiyattaki marj ve hedefe inmek için gereken maliyet düşüşü gösterilir. Kur ±%5 ve kumaş fiyatı ±%10
duyarlılığı hesaplanır. İnternete bağlanmaz.

Kullanım:
    python main.py                                                     # örnek: 3 model, örnek kurlar
    python main.py --modeller modeller.xlsx --kalemler maliyet_kalemleri.xlsx --kur USD=41,20 EUR=48,10
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
ORNEK_KUR = {"USD": Decimal("41.20"), "EUR": Decimal("48.10")}
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
GRUPLAR = ["Kumaş", "Aksesuar", "Fason İşlem", "CM", "Diğer", "Sabit"]
GRUP_ADLARI = {"Kumaş": ("kumas", "fabric", "ana kumas", "garni"), "Aksesuar": ("aksesuar", "trim", "trims", "malzeme"),
               "Fason İşlem": ("fason islem", "fason", "baski", "nakis", "yikama", "islem", "embellishment"),
               "CM": ("cm", "iscilik", "kesim dikim", "cut make"), "Diğer": ("diger", "other", "adet basi gider"),
               "Sabit": ("sabit", "siparis basi", "sabit gider", "fixed")}
PARA = {"TL": ("tl", "try", "trl"), "USD": ("usd", "dolar"), "EUR": ("eur", "euro", "avro"), "GBP": ("gbp", "sterlin")}
MODEL_SUTUNLARI = {"model": ("model", "style", "model kodu"), "musteri": ("musteri", "buyer"), "aciklama": ("aciklama", "tanim", "urun"),
                   "adet": ("siparis adedi", "adet", "miktar"), "para": ("teklif para birimi", "para birimi", "doviz"),
                   "verim": ("verimlilik", "verimlilik yuzde", "hat verimliligi"), "gg": ("genel gider", "genel gider yuzde", "overhead"),
                   "marj": ("kar marji", "kar marji yuzde", "marj"), "komisyon": ("komisyon", "komisyon yuzde", "acente komisyonu"),
                   "hedef": ("hedef fiyat", "musteri hedef fiyati", "target price")}
KALEM_SUTUNLARI = {"model": ("model", "style"), "grup": ("grup", "kalem grubu", "tur"), "kalem": ("kalem", "aciklama", "malzeme"),
                   "miktar": ("miktar", "tuketim", "sam", "miktar tuketim"), "birim": ("birim",), "fire": ("fire", "fire yuzde", "fire orani"),
                   "fiyat": ("birim fiyat", "fiyat"), "para": ("para birimi", "doviz", "para")}
K2 = Decimal("0.01")
K4 = Decimal("0.0001")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def sayi(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("%", "").replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuzde(x) -> Decimal:
    d = sayi(x)
    if d is None:
        return Decimal(0)
    return d / 100 if d >= 1 else d          # 8 veya 0,08


def para_bul(x) -> str | None:
    k = katla(x) or "tl"
    return next((p for p, adlar in PARA.items() if k in adlar), k.upper())


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


def kur_coz(degerler: list[str] | None) -> dict[str, Decimal]:
    kur = {}
    for d in degerler or []:
        m = re.fullmatch(r"\s*([A-Za-z]{3})\s*=\s*([\d.,]+)\s*", d)
        if not m or sayi(m.group(2)) is None:
            raise ValueError(f"Kur '{d}' okunamadı; örnek: USD=41,20")
        kur[m.group(1).upper()] = sayi(m.group(2))
    return kur


def grup_bul(x) -> str | None:
    k = katla(x)
    return next((g for g, adlar in GRUP_ADLARI.items() if k in adlar or any(k.startswith(a + " ") for a in adlar)), None)


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Model:
    kod: str
    musteri: str
    aciklama: str
    adet: int
    para: str
    verim: Decimal
    gg: Decimal
    marj: Decimal
    komisyon: Decimal
    hedef: Decimal | None


@dataclass
class Kalem:
    satir: int
    model: str
    grup: str
    ad: str
    miktar: Decimal
    birim: str
    fire: Decimal
    fiyat: Decimal | None
    para: str
    tutar: Decimal | None = None          # teklif para biriminde, adet başı
    not_: str = ""


@dataclass
class Sonuc:
    model: Model
    kalemler: list
    gruplar: dict = field(default_factory=dict)
    uretim: Decimal = Decimal(0)
    genel_gider: Decimal = Decimal(0)
    toplam: Decimal = Decimal(0)
    fob: Decimal = Decimal(0)
    hedef_marj: Decimal | None = None
    hedef_dusus: Decimal | None = None
    duyarlilik: list = field(default_factory=list)


def modelleri_oku(yol: Path) -> dict[str, Model]:
    sonuc = {}
    for r in kayitlar(yol, MODEL_SUTUNLARI, ("model",)):
        kod = str(r.get("model") or "").strip()
        if not kod:
            continue
        verim = yuzde(r.get("verim")) or Decimal(1)
        sonuc[kod] = Model(kod, str(r.get("musteri") or "").strip(), str(r.get("aciklama") or "").strip(), int(sayi(r.get("adet")) or 0),
                           para_bul(r.get("para") or "USD"), verim, yuzde(r.get("gg")), yuzde(r.get("marj")), yuzde(r.get("komisyon")), sayi(r.get("hedef")))
    return sonuc


def kalemleri_oku(yol: Path) -> tuple[list[Kalem], list[str]]:
    sonuc, hatalar = [], []
    for r in kayitlar(yol, KALEM_SUTUNLARI, ("model", "grup", "miktar", "fiyat")):
        if not r.get("model"):
            continue
        grup = grup_bul(r.get("grup"))
        miktar = sayi(r.get("miktar"))
        if grup is None or miktar is None:
            hatalar.append(f"Satır {r['_satir']}: " + (f"grup tanınmadı ('{r.get('grup')}')" if grup is None else "miktar okunamadı"))
            continue
        sonuc.append(Kalem(r["_satir"], str(r["model"]).strip(), grup, str(r.get("kalem") or "").strip(), miktar, str(r.get("birim") or "").strip(),
                           yuzde(r.get("fire")), sayi(r.get("fiyat")), para_bul(r.get("para"))))
    return sonuc, hatalar


# ----------------------------------------------------------------------------
# Hesap
# ----------------------------------------------------------------------------

def cevir(tutar: Decimal, kaynak: str, hedef: str, kur: dict[str, Decimal]) -> Decimal | None:
    """kur: 1 birim döviz = x TL."""
    if kaynak == hedef:
        return tutar
    tl = tutar if kaynak == "TL" else (tutar * kur[kaynak] if kaynak in kur else None)
    if tl is None:
        return None
    return tl if hedef == "TL" else (tl / kur[hedef] if hedef in kur else None)


def maliyetle(m: Model, kalemler: list[Kalem], kur: dict[str, Decimal], uyarilar: list, kumas_carpan=Decimal(1), kur_carpan=Decimal(1),
              kayit: bool = True) -> Sonuc:
    """kur_carpan: teklif para biriminin TL karşılığı çarpanı (0,95 = döviz %5 ucuzladı)."""
    kur2 = dict(kur)
    if m.para in kur2:
        kur2[m.para] = kur[m.para] * kur_carpan
    s = Sonuc(m, kalemler, {g: Decimal(0) for g in GRUPLAR})

    def uyar(onem, tur, aciklama):
        if kayit:
            uyarilar.append({"onem": onem, "tur": tur, "model": m.kod, "aciklama": aciklama})

    for k in kalemler:
        if k.fiyat is None:
            k.tutar, k.not_ = None, "Fiyat yok — hesaba katılmadı"
            uyar("Yüksek", "Fiyat eksik", f"{k.grup} / {k.ad}: birim fiyat boş; maliyet bu kalem olmadan hesaplandı")
            continue
        if k.grup == "CM":
            birim_maliyet = k.miktar * k.fiyat / m.verim
        elif k.grup == "Sabit":
            if not m.adet:
                k.tutar, k.not_ = None, "Sipariş adedi yok"
                uyar("Yüksek", "Sipariş adedi yok", f"{k.ad}: sabit gider adede dağıtılamadı")
                continue
            birim_maliyet = k.miktar * k.fiyat / m.adet
        else:
            birim_maliyet = k.miktar * (1 + k.fire) * k.fiyat * (kumas_carpan if k.grup == "Kumaş" else 1)
        t = cevir(birim_maliyet, k.para, m.para, kur2)
        if t is None:
            k.tutar, k.not_ = None, f"{k.para} kuru yok — hesaba katılmadı"
            uyar("Yüksek", "Kur eksik", f"{k.grup} / {k.ad}: {k.para} kuru verilmedi (--kur {k.para}=...); kalem hesaba katılmadı")
            continue
        k.tutar = t
        s.gruplar[k.grup] += t
        if kayit and k.grup == "Kumaş" and k.fire == 0:
            uyar("Bilgi", "Kumaş firesi 0", f"{k.ad}: fire oranı girilmemiş; pastal / kesim firesini kontrol edin")
    s.uretim = sum(s.gruplar.values(), Decimal(0))
    s.genel_gider = s.uretim * m.gg
    s.toplam = s.uretim + s.genel_gider
    payda = 1 - m.marj - m.komisyon
    if payda <= 0:
        raise ValueError(f"{m.kod}: kâr marjı + komisyon %100'den küçük olmalı")
    s.fob = s.toplam / payda
    if m.hedef:
        s.hedef_marj = (m.hedef * (1 - m.komisyon) - s.toplam) / m.hedef
        izin = m.hedef * payda
        s.hedef_dusus = max(s.toplam - izin, Decimal(0))
    return s


def hesapla(modeller: dict[str, Model], kalemler: list[Kalem], kur: dict[str, Decimal]) -> dict:
    uyarilar = []
    gruplu = defaultdict(list)
    for k in kalemler:
        gruplu[k.model].append(k)
    for kod in sorted(set(gruplu) - set(modeller)):
        uyarilar.append({"onem": "Orta", "tur": "Model tanımsız", "model": kod, "aciklama": f"{len(gruplu[kod])} kalem var ama modeller tablosunda yok"})
    sonuclar = []
    for kod, m in modeller.items():
        lst = gruplu.get(kod, [])
        if not lst:
            uyarilar.append({"onem": "Yüksek", "tur": "Kalem yok", "model": kod, "aciklama": "Modelin maliyet kalemi yok"})
            continue
        if m.para != "TL" and m.para not in kur:
            uyarilar.append({"onem": "Yüksek", "tur": "Kur eksik", "model": kod, "aciklama": f"Teklif para birimi {m.para} için kur yok"})
            continue
        s = maliyetle(m, lst, kur, uyarilar)
        if not any(k.grup == "Kumaş" for k in lst):
            uyarilar.append({"onem": "Orta", "tur": "Kumaş kalemi yok", "model": kod, "aciklama": "Maliyet föyünde kumaş kalemi yok"})
        if not any(k.grup == "CM" for k in lst):
            uyarilar.append({"onem": "Orta", "tur": "CM kalemi yok", "model": kod, "aciklama": "Kesim-dikim (CM) kalemi yok"})
        if m.hedef and s.hedef_marj is not None and s.hedef_marj < m.marj:
            uyarilar.append({"onem": "Yüksek" if s.hedef_marj < 0 else "Orta", "tur": "Hedef fiyat altında", "model": kod,
                             "aciklama": f"Hedef {fmt(m.hedef)} {m.para} fiyatında marj %{fmt(s.hedef_marj * 100, 1)} (istenen %{fmt(m.marj * 100, 1)}); "
                                         f"hedefe inmek için adet başı {fmt(s.hedef_dusus)} {m.para} maliyet düşüşü gerekir"})
        if m.para != "TL":
            for kc in (Decimal("0.95"), Decimal("1.05")):
                d = maliyetle(m, lst, kur, uyarilar, kur_carpan=kc, kayit=False)
                s.duyarlilik.append((f"{m.para}/TL kuru %{'-' if kc < 1 else '+'}5", d.toplam, d.fob))
        for kc in (Decimal("0.90"), Decimal("1.10")):
            d = maliyetle(m, lst, kur, uyarilar, kumas_carpan=kc, kayit=False)
            s.duyarlilik.append((f"Kumaş fiyatı %{'-' if kc < 1 else '+'}10", d.toplam, d.fob))
        maliyetle(m, lst, kur, uyarilar, kayit=False)       # kalem tutarlarını temel senaryoya geri yaz
        sonuclar.append(s)
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    uyarilar.sort(key=lambda u: (sira[u["onem"]], u["model"]))
    return {"sonuclar": sonuclar, "uyarilar": uyarilar}


def fmt(x: Decimal, n: int = 2) -> str:
    q = Decimal(1).scaleb(-n)
    s = f"{x.quantize(q, ROUND_HALF_UP):,.{n}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KONTROL = PatternFill("solid", fgColor="FFF4CE")
UST = Alignment(vertical="top", wrap_text=True)
F4 = "#,##0.0000"
F2 = "#,##0.00"


def _baslik(ws, basliklar, genislik=None):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik or (), 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def _f(x, q=K4):
    return None if x is None else float(x.quantize(q, ROUND_HALF_UP))


def rapor_yaz(cikti: Path, s: dict, kur: dict, hatalar: list[str]) -> None:
    wb = Workbook()
    o = wb.active
    o.title = "Özet"
    _baslik(o, ["Model", "Müşteri", "Açıklama", "Adet", "Para", "Kumaş", "Aksesuar", "Fason İşlem", "CM", "Diğer", "Sabit (adet başı)",
                "Üretim Maliyeti", "Genel Gider", "Toplam Maliyet", "Marj %", "Komisyon %", "FOB Teklif", "Hedef Fiyat", "Hedefte Marj",
                "Gereken Düşüş", "Karar / Not"], (8, 16, 30, 8, 6) + (10,) * 9 + (8, 9, 10, 10, 10, 10, 26))
    o.freeze_panes = "B2"
    for x in s["sonuclar"]:
        m = x.model
        o.append([m.kod, m.musteri, m.aciklama, m.adet, m.para] + [_f(x.gruplar[g]) for g in GRUPLAR] +
                 [_f(x.uretim), _f(x.genel_gider), _f(x.toplam), float(m.marj), float(m.komisyon), _f(x.fob, K2), _f(m.hedef, K2) if m.hedef else None,
                  float(x.hedef_marj) if x.hedef_marj is not None else None, _f(x.hedef_dusus), ""])
        r = o.max_row
        for j in range(6, 15):
            o.cell(r, j).number_format = F4
        for j in (15, 16, 19):
            o.cell(r, j).number_format = "0.0%"
        o.cell(r, 17).number_format = F2
        o.cell(r, 17).font = Font(bold=True)
        o.cell(r, 20).number_format = F4
        if x.hedef_marj is not None and x.hedef_marj < m.marj:
            o.cell(r, 19).fill = PatternFill("solid", fgColor=RENK["Yüksek" if x.hedef_marj < 0 else "Orta"])
        o.cell(r, 21).fill = KONTROL
    o.append([])
    o.append(["Kurlar (1 birim = TL)", ", ".join(f"{p} = {fmt(v, 4)}" for p, v in kur.items()) or "—"])
    o.append(["Formül", "FOB = toplam maliyet ÷ (1 − marj − komisyon); CM = SAM × dakika maliyeti ÷ verimlilik"])
    for h in hatalar:
        o.append(["Okunamayan satır", h])

    for x in s["sonuclar"]:
        m = x.model
        ws = wb.create_sheet(f"Föy {m.kod}"[:31])
        ws.append([f"Maliyet föyü · {m.kod} · {m.musteri}"])
        ws.cell(1, 1).font = Font(bold=True, size=12)
        ws.append([m.aciklama, f"Sipariş {m.adet:,} adet".replace(",", "."), f"Teklif para birimi {m.para}", f"Verimlilik %{fmt(m.verim * 100, 0)}"])
        ws.append([])
        _baslik(ws, ["Grup", "Kalem", "Miktar", "Birim", "Fire %", "Birim Fiyat", "Para", f"Adet Başı ({m.para})", "Pay %", "Not"],
                (12, 38, 9, 9, 7, 11, 6, 14, 8, 30))
        for g in GRUPLAR:
            for k in [k for k in x.kalemler if k.grup == g]:
                ws.append([k.grup, k.ad, float(k.miktar), k.birim, float(k.fire) if k.grup not in ("CM", "Sabit", "Diğer") else None,
                           float(k.fiyat) if k.fiyat is not None else None, k.para, _f(k.tutar),
                           float(k.tutar / x.toplam) if k.tutar is not None and x.toplam else None, k.not_])
                r = ws.max_row
                ws.cell(r, 5).number_format = "0%"
                ws.cell(r, 8).number_format = F4
                ws.cell(r, 9).number_format = "0.0%"
                if k.not_:
                    ws.cell(r, 10).fill = PatternFill("solid", fgColor=RENK["Yüksek"])
            if x.gruplar[g]:
                ws.append(["", f"{g} toplamı", None, None, None, None, None, _f(x.gruplar[g])])
                ws.cell(ws.max_row, 2).font = ws.cell(ws.max_row, 8).font = Font(bold=True)
                ws.cell(ws.max_row, 8).number_format = F4
        ws.append([])
        for etiket, deger, bicim in (("Üretim maliyeti", x.uretim, F4), (f"Genel gider (%{fmt(m.gg * 100, 1)})", x.genel_gider, F4),
                                     ("Toplam maliyet", x.toplam, F4), (f"FOB teklif (marj %{fmt(m.marj * 100, 1)}, komisyon %{fmt(m.komisyon * 100, 1)})", x.fob, F2)):
            ws.append(["", etiket, None, None, None, None, None, _f(deger)])
            ws.cell(ws.max_row, 2).font = ws.cell(ws.max_row, 8).font = Font(bold=True)
            ws.cell(ws.max_row, 8).number_format = bicim
        if m.hedef:
            ws.append(["", "Müşteri hedef fiyatı", None, None, None, None, None, _f(m.hedef, K2)])
            ws.append(["", "Hedef fiyatta marj", None, None, None, None, None, float(x.hedef_marj)])
            ws.cell(ws.max_row, 8).number_format = "0.0%"
        ws.freeze_panes = "A5"

    dy = wb.create_sheet("Duyarlılık")
    _baslik(dy, ["Model", "Senaryo", "Toplam Maliyet", "FOB Teklif", "Temel FOB'a Göre"], (8, 24, 14, 12, 14))
    for x in s["sonuclar"]:
        dy.append([x.model.kod, "Temel", _f(x.toplam), _f(x.fob, K2), 0])
        for ad, toplam, fob in x.duyarlilik:
            dy.append([x.model.kod, ad, _f(toplam), _f(fob, K2), float((fob - x.fob) / x.fob)])
        for r in dy.iter_rows(min_row=dy.max_row - len(x.duyarlilik), max_row=dy.max_row):
            r[2].number_format, r[3].number_format, r[4].number_format = F4, F2, "+0.0%;-0.0%;0.0%"

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Model", "Açıklama", "İnceleme"], (9, 22, 8, 90, 24))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["model"], u["aciklama"], ""])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).alignment = UST
        uy.cell(uy.max_row, 5).fill = KONTROL
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(model_yolu: Path, kalem_yolu: Path, cikti: Path, kur: dict[str, Decimal]) -> dict:
    modeller = modelleri_oku(model_yolu)
    kalemler, hatalar = kalemleri_oku(kalem_yolu)
    s = hesapla(modeller, kalemler, kur)
    rapor_yaz(cikti, s, kur, hatalar)
    return {**s, "modeller": modeller, "hatalar": hatalar}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Kumaş, aksesuar, fason işlem, CM ve genel giderlerden hazır giyim FOB maliyeti ve teklif fiyatı hesaplar.")
    p.add_argument("--modeller", type=Path, default=ORNEK / "modeller.csv", help="Model, sipariş adedi, teklif para birimi, verimlilik, GG / marj / komisyon %%, hedef fiyat")
    p.add_argument("--kalemler", type=Path, default=ORNEK / "maliyet_kalemleri.csv", help="Model, Grup, Kalem, Miktar, Birim, Fire %%, Birim Fiyat, Para Birimi")
    p.add_argument("--kur", nargs="*", help="1 birim döviz = TL, ör. USD=41,20 EUR=48,10 (örnek veride örnek kurlar kullanılır)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "maliyet_foyu.xlsx")
    a = p.parse_args(argv)
    for y in (a.modeller, a.kalemler):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        kur = kur_coz(a.kur)
        if not kur and a.modeller == ORNEK / "modeller.csv":
            kur = dict(ORNEK_KUR)
            print("[i] Örnek kurlar kullanılıyor (USD=41,20 EUR=48,10); kendi verinizde --kur verin.")
        s = calistir(a.modeller, a.kalemler, a.cikti, kur)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    for x in s["sonuclar"]:
        m = x.model
        print(f"[OK] {m.kod}: toplam maliyet {fmt(x.toplam, 4)} {m.para} · FOB teklif {fmt(x.fob)} {m.para}"
              + (f" · hedef {fmt(m.hedef)} → marj %{fmt(x.hedef_marj * 100, 1)}" if m.hedef else ""))
    for u in s["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['model']} {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
