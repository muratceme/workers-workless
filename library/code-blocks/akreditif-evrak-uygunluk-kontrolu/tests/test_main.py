import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
IBRAZ = date(2026, 10, 20)


def akr(**kw):
    a = {"no": "LC1", "pb": "USD", "tutar": D("10000"), "tutar_tol": None, "miktar": D("1000"), "birim": "kg", "miktar_tol": None, "son_yukleme": date(2026, 10, 10),
         "vade": date(2026, 11, 6), "ibraz_gun": None, "kismi": "izinli", "lehtar": "Örnek A.Ş.", "amir": "Sample GmbH", "mal": "COTTON YARN NE 30/1", "teslim": "CIF",
         "yukleme_limani": "Mersin", "bosaltma_limani": "Hamburg", "sigorta_oran": None, "mense": ""}
    return a | kw


def evrak(**kw):
    e = {"fatura": {"_ad": "Fatura", "tutar": "10.000,00", "para birimi": "USD", "miktar": "1000", "mal tanimi": "Cotton yarn Ne 30/1", "duzenleyen": "ÖRNEK A.Ş. Mersin",
                    "alici": "Sample GmbH", "asil adet": 1},
         "konsimento": {"_ad": "Konşimento", "yukleme tarihi": "05.10.2026", "asil adet": 3}}
    for k, v in kw.items():
        e.setdefault(k, {"_ad": k}).update(v)
    return e


def konular(r, onem="Yüksek"):
    return [x["konu"] for x in r["bulgular"] if x["onem"] == onem]


class SureTesti(unittest.TestCase):
    def test_temiz_set(self):
        r = main.kontrol(akr(), [], evrak(), IBRAZ)
        self.assertEqual(konular(r), [])
        self.assertEqual(r["son_ibraz"], date(2026, 10, 26))                       # 05.10 + 21 gün

    def test_gec_ibraz_ve_vade(self):
        self.assertEqual(konular(main.kontrol(akr(), [], evrak(), date(2026, 10, 27))), ["Geç ibraz"])
        self.assertEqual(konular(main.kontrol(akr(vade=date(2026, 10, 19)), [], evrak(), IBRAZ)), ["Akreditifin vadesi geçmiş"])
        self.assertEqual(konular(main.kontrol(akr(ibraz_gun=D(10)), [], evrak(), date(2026, 10, 16))), ["Geç ibraz"])

    def test_hafta_sonu_uzar(self):
        r = main.kontrol(akr(vade=date(2026, 10, 24)), [], evrak(), date(2026, 10, 26))      # Cumartesi → Pazartesi
        self.assertEqual(konular(r), [])
        self.assertEqual(len(konular(r, "Bilgi")), 1)
        erken = evrak(konsimento={"yukleme tarihi": "03.10.2026"})                            # 21. gün Cumartesi 24.10 → 26.10
        self.assertEqual(konular(main.kontrol(akr(), [], erken, date(2026, 10, 26))), [])
        self.assertEqual(konular(main.kontrol(akr(), [], erken, date(2026, 10, 27))), ["Geç ibraz"])

    def test_gec_yukleme(self):
        self.assertEqual(konular(main.kontrol(akr(), [], evrak(konsimento={"yukleme tarihi": "11.10.2026"}), IBRAZ)), ["Geç yükleme"])


class EvrakTesti(unittest.TestCase):
    def test_tutar_ve_tolerans(self):
        self.assertEqual(konular(main.kontrol(akr(), [], evrak(fatura={"tutar": "10.000,01"}), IBRAZ)), ["Fatura tutarı akreditif tutarını aşıyor"])
        self.assertEqual(konular(main.kontrol(akr(tutar_tol="about"), [], evrak(fatura={"tutar": "11.000,00"}), IBRAZ)), [])
        self.assertEqual(len(konular(main.kontrol(akr(tutar_tol="5"), [], evrak(fatura={"tutar": "10.500,01"}), IBRAZ))), 1)

    def test_miktar(self):
        def k(a, m):
            return konular(main.kontrol(a, [], evrak(fatura={"miktar": m}), IBRAZ))
        self.assertEqual(k(akr(), "1050"), [])                                     # kg: ±%5 (md. 30(b))
        self.assertEqual(k(akr(), "1051"), ["Miktar akreditiftekinden fazla"])
        self.assertEqual(k(akr(birim="adet"), "1001"), ["Miktar akreditiftekinden fazla"])   # adet: tolerans yok
        self.assertEqual(k(akr(kismi="yasak"), "940"), ["Kısmi yükleme yasak, miktar eksik"])
        self.assertEqual(k(akr(), "940"), [])                                      # kısmi serbest → yalnız bilgi

    def test_fatura_taraflar_ve_mal(self):
        r = main.kontrol(akr(), [], evrak(fatura={"duzenleyen": "Başka Ltd.", "alici": "Other GmbH", "mal tanimi": "Cotton yarn", "para birimi": "EUR"}), IBRAZ)
        self.assertEqual({x["madde"] for x in r["bulgular"]}, {"md. 18(a)(i)", "md. 18(a)(ii)", "md. 18(a)(iii)", "md. 18(c)"})

    def test_konsimento(self):
        r = main.kontrol(akr(), [], evrak(konsimento={"serh": "3 cartons wet", "yukleme kaydi": "Received for shipment", "duzenlenen asil": 3, "asil adet": 2,
                                                       "bosaltma limani": "Rotterdam"}), IBRAZ)
        self.assertEqual({x["madde"] for x in r["bulgular"]}, {"md. 27", "md. 20(a)(ii)", "md. 20(a)(iv)", "md. 20(a)(iii)"})

    def test_sigorta(self):
        def f(**kw):
            return konular(main.kontrol(akr(), [], evrak(sigorta={"sigorta tutari": "11.000,00", "para birimi": "USD", "tarih": "05.10.2026"} | kw), IBRAZ))
        self.assertEqual(f(), [])
        self.assertEqual(f(**{"sigorta tutari": "10.999,99"}), ["Sigorta tutarı yetersiz"])
        self.assertEqual(f(**{"para birimi": "EUR"}), ["Sigorta para birimi akreditiften farklı"])
        self.assertEqual(f(tarih="06.10.2026"), ["Sigorta tarihi yüklemeden sonra"])
        self.assertEqual(f(tarih="06.10.2026", **{"teminat baslangici": "04.10.2026"}), [])

    def test_istenen_ve_capraz(self):
        ist = [{"ad": "Konşimento", "tur": "konsimento", "asil": 3, "kopya": 0, "sart": ""}, {"ad": "Menşe Şahadetnamesi", "tur": "mense", "asil": 1, "kopya": 0, "sart": ""},
               {"ad": "Ticari Fatura", "tur": "fatura", "asil": 2, "kopya": 0, "sart": ""}]
        r = main.kontrol(akr(), ist, evrak(fatura={"koli": "100", "tarih": "21.10.2026"}, ceki={"koli": "101"}), IBRAZ)
        self.assertEqual([x["durum"] for x in r["liste"]], ["Tamam", "Eksik", "Adet eksik"])
        self.assertIn("Koli sayısı evraklar arasında çelişiyor", konular(r))
        self.assertIn("Evrak tarihi ibraz tarihinden sonra", konular(r))
        self.assertEqual(main.evrak_turu("Sigorta Poliçesi"), "sigorta")
        self.assertEqual(main.evrak_turu("Poliçe (Draft)"), "police")


class UctanUca(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "a.xlsx"
            r = main.calistir(ORNEK / "akreditif.csv", ORNEK / "istenen_evraklar.csv", ORNEK / "evraklar.csv", date(2026, 11, 2), cikti)
            self.assertEqual(sorted(konular(r)), sorted(["Geç yükleme", "Asıl / kopya sayısı eksik", "Mal tanımı akreditiftekiyle uyuşmuyor", "Sigorta tutarı yetersiz",
                                                          "Sigorta tarihi yüklemeden sonra", "Brüt ağırlık evraklar arasında çelişiyor"]))
            self.assertEqual((r["vade_fiili"], r["son_ibraz"]), (date(2026, 11, 2), date(2026, 11, 2)))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Rezervler", "Evrak Listesi", "Çapraz Kontrol"])

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "a.xlsx")]), 0)
            self.assertEqual(main.main(["--ibraz", "yarın"]), 2)
            self.assertEqual(main.main(["--akreditif", str(Path(tmp) / "yok.csv")]), 1)
            kotu = Path(tmp) / "k.csv"
            kotu.write_text("Alan;Değer\nTutar;100\n", encoding="utf-8")
            self.assertEqual(main.main(["--akreditif", str(kotu)]), 1)


if __name__ == "__main__":
    unittest.main()
