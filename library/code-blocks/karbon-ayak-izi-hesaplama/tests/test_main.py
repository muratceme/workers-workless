import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
FAK = {"dogal gaz": {"kaynak": "Doğal gaz", "birim": "m³", "faktor": D("2.00"), "kapsam": 1, "kaynakca": "x"},
       "motorin": {"kaynak": "Motorin", "birim": "L", "faktor": D("2.68"), "kapsam": 1, "kaynakca": "x"},
       "elektrik": {"kaynak": "Elektrik", "birim": "kWh", "faktor": D("0.44"), "kapsam": 2, "kaynakca": "x"},
       "kumas": {"kaynak": "Kumaş", "birim": "kg", "faktor": D(5), "kapsam": 3, "kaynakca": "x"}}


def t(kaynak, miktar, birim, donem=(2026, 7), tesis="F1"):
    return {"donem": donem, "tesis": tesis, "kaynak": kaynak, "miktar": D(miktar), "birim": birim, "kim": kaynak}


class HesapTesti(unittest.TestCase):
    def test_kapsamlar_ve_urun_basina(self):
        r = main.hesapla([t("Doğal gaz", 10000, "m3"), t("Motorin", 1000, "lt"), t("ELEKTRİK", 100000, "kWh"), t("Kumaş", 100, "kg")], FAK, {(2026, 7): D(50000)})
        self.assertEqual((r["kapsam"], r["toplam"], r["ton"], r["urun_basina"]), ({1: D("22680.00"), 2: D("44000.00"), 3: D(500)}, D("67180.00"), D("67.18"), D("1.3436")))
        self.assertEqual([(x["kaynak"], x["emisyon"], x["pay"]) for x in r["kaynak"]],
                         [("Elektrik", D("44000.00"), D("65.50")), ("Doğal gaz", D("20000.00"), D("29.77")), ("Motorin", D("2680.00"), D("3.99")), ("Kumaş", D(500), D("0.74"))])

    def test_donem_ve_tesis(self):
        l = [t("Elektrik", 1000, "kWh"), t("Elektrik", 2000, "kWh", (2026, 8)), t("Doğal gaz", 100, "m³", (2026, 8), "F2")]
        r = main.hesapla(l, FAK, {(2026, 7): D(100), (2026, 9): D(5)})
        self.assertEqual([(x["donem"], x["emisyon"], x["uretim"], x["urun_basina"]) for x in r["donemler"]], [((2026, 7), D("440.00"), D(100), D("4.4000")), ((2026, 8), D("1080.00"), None, None)])
        self.assertEqual((r["uretim"], r["urun_basina"]), (D(100), D("4.4000")))                       # yalnız üretimi bilinen dönemler
        self.assertEqual([(x["tesis"], x["emisyon"]) for x in r["tesis"]], [("F1", D("1320.00")), ("F2", D("200.00"))])
        self.assertEqual([u["tur"] for u in r["uyarilar"]], ["Üretimi olan dönemin tüketimi yok"])
        self.assertIsNone(main.hesapla(l, FAK)["urun_basina"])

    def test_eslesmeyen_satirlar(self):
        r = main.hesapla([t("Elektrik", 1000, "kWh"), t("Elektrik", 5, "MWh"), t("LPG", 500, "kg"), t("LPG", 1, "kg"), t("Motorin", 10, "")], FAK)
        self.assertEqual((len(r["satirlar"]), r["toplam"]), (2, D("466.80")))                             # birimi boş satır faktör birimiyle kabul edilir
        self.assertEqual(sorted((u["tur"], u["aciklama"][:1]) for u in r["uyarilar"]), [("Birim faktörle uyuşmuyor", "T"), ("Kaynağın emisyon faktörü yok", "2")])
        with self.assertRaises(ValueError):
            main.hesapla([t("LPG", 1, "kg")], FAK)
        self.assertEqual([main.birim_katla(x) for x in ("m³", "M3", "Lt", "litre", "kWh")], ["m3", "m3", "l", "l", "kwh"])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "k.xlsx"
            r = main.calistir(ORNEK / "tuketimler.csv", ORNEK / "emisyon_faktorleri.csv", cikti, ORNEK / "uretim.csv")
            self.assertEqual((r["kapsam"][1], r["kapsam"][2], r["toplam"], r["ton"], r["uretim"], r["urun_basina"]), (D("68308.00"), D("132000.00"), D("200308.00"), D("200.31"), D(150000), D("1.3354")))
            self.assertEqual([(x["emisyon"], x["urun_basina"]) for x in r["donemler"]], [(D("66680.00"), D("1.3336")), (D("58816.00"), D("1.4704")), (D("74812.00"), D("1.2469"))])
            self.assertEqual([(x["kaynak"], x["pay"]) for x in r["kaynak"]], [("Elektrik", D("65.90")), ("Doğal gaz", D("29.95")), ("Motorin", D("4.15"))])
            self.assertEqual([u["tur"] for u in r["uyarilar"]], ["Birim faktörle uyuşmuyor", "Kaynağın emisyon faktörü yok"])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Dönem Bazında", "Kaynak Bazında", "Tesis Bazında", "Faktörler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            tk, f, u = (Path(tmp) / x for x in ("t.csv", "f.csv", "u.csv"))
            f.write_text("Kaynak;Birim;Emisyon Faktörü;Kapsam;Kaynakça\nElektrik;kWh;0,5;2;\nGaz;m3;çok;1;x\nKömür;kg;2,4;4;x\n", encoding="utf-8")
            tk.write_text("Dönem;Kaynak;Miktar;Birim\n2026-07;Elektrik;1.000;kWh\ntemmuz;Elektrik;5;kWh\n07.2026;Elektrik;yok;kWh\n", encoding="utf-8")
            u.write_text("Dönem;Üretim Miktarı\n07.2026;200\n07.2026;300\n08.2026;sıfır\n", encoding="utf-8")
            tuk, fak, ure, uy = main.oku(tk, f, u)
            self.assertEqual((list(fak), [(x["donem"], x["miktar"]) for x in tuk], ure), (["elektrik"], [((2026, 7), D(1000))], {(2026, 7): D(500)}))
            self.assertEqual(sorted(x["tur"] for x in uy), ["Dönem okunamadı", "Emisyon faktörü okunamadı", "Faktörün kaynağı yazılmamış", "Kapsam okunamadı", "Tüketim miktarı okunamadı", "Üretim satırı okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--tuketimler", str(tk), "--faktorler", str(f), "--cikti", c]
            self.assertEqual(main.main(arg + ["--uretim", str(u)]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(["--tuketimler", str(Path(tmp) / "yok.csv")]), 1)
            f.write_text("Kaynak;Emisyon Faktörü;Kapsam\nGaz;x;1\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
