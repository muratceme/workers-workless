"""
Karbon Ayak İzi Hesaplama — Gerisi Bende kod bloğu
Sürdürülebilirlik ve Sosyal Uygunluk › Sosyal Uygunluk Uzmanı

Enerji, yakıt ve malzeme tüketim verisinden kapsam bazında emisyonu ve ürün başına karbon değerini hesaplar.
  Emisyon (kg CO2e)  = tüketim miktarı × emisyon faktörü (kg CO2e ÷ birim)
  Kapsam             = faktör dosyasındaki kapsam (1 = tesiste yakılan yakıt, 2 = satın alınan elektrik / ısı, 3 = diğer dolaylı)
  Ürün başına        = dönemin toplam emisyonu ÷ dönemin üretim miktarı
Emisyon faktörleri koda gömülü değildir; kaynağını belirterek siz verirsiniz. Tüketim birimi faktör birimiyle aynı olmalıdır (birim çevrilmez).
Yöntemsel bir envanter değildir: sınır tanımı, kaçak emisyonlar (soğutucu gaz), proses emisyonları ve doğrulama bu hesabın dışındadır.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


TUK_SUTUN = {"donem": ("donem", "ay", "tarih"), "tesis": ("tesis", "fabrika", "lokasyon"), "kaynak": ("kaynak", "enerji kaynagi", "yakit", "malzeme", "kalem"),
             "miktar": ("miktar", "tuketim"), "birim": ("birim",)}
FAK_SUTUN = {"kaynak": ("kaynak", "enerji kaynagi", "yakit", "malzeme", "kalem"), "birim": ("birim",), "faktor": ("emisyon faktoru", "emisyon faktoru kg co2e birim", "faktor", "kg co2e birim", "ef"),
             "kapsam": ("kapsam", "scope"), "kaynakca": ("kaynakca", "referans", "faktor kaynagi")}
URE_SUTUN = {"donem": ("donem", "ay", "tarih"), "tesis": ("tesis", "fabrika", "lokasyon"), "miktar": ("uretim miktari", "uretim adedi", "uretim", "miktar"), "birim": ("birim",)}
Z = Decimal(0)


def donem(x) -> tuple[int, int] | None:
    if isinstance(x, (date, datetime)):
        return x.year, x.month
    s = metin(x)
    m = re.fullmatch(r"(\d{4})[-./ ](\d{1,2})", s)
    if m:
        y, a = int(m[1]), int(m[2])
    else:
        m = re.fullmatch(r"(\d{1,2})[-./ ](\d{4})", s)
        if not m:
            t = tarih(s)
            return (t.year, t.month) if t else None
        y, a = int(m[2]), int(m[1])
    return (y, a) if 1 <= a <= 12 else None


def donem_ad(d) -> str:
    return f"{d[1]:02d}.{d[0]}" if d else "—"


def birim_katla(x) -> str:
    return metin(x).lower().replace("³", "3").replace("²", "2").replace(" ", "").replace("lt", "l").replace("litre", "l")


def oku(tuk_yol: Path, fak_yol: Path, ure_yol: Path | None) -> tuple[list[dict], dict, dict, list[dict]]:
    uy, fak, tuk, ure = [], {}, [], defaultdict(lambda: Z)
    for r in kayitlar(fak_yol, FAK_SUTUN, ("kaynak", "faktor")):
        k = metin(r["kaynak"])
        if not k:
            continue
        f, kp = para(r["faktor"]), para(r["kapsam"])
        if f is None or f < 0:
            uy.append(uyari("Yüksek", "Emisyon faktörü okunamadı", k, f"{fak_yol.name} satır {r['_satir']}: '{metin(r['faktor'])}'. Kaynak hesaba alınamaz."))
            continue
        if kp is None or int(kp) not in (1, 2, 3):
            uy.append(uyari("Yüksek", "Kapsam okunamadı", k, f"{fak_yol.name} satır {r['_satir']}: '{metin(r['kapsam'])}'. 1, 2 veya 3 olmalı; kaynak hesaba alınamaz."))
            continue
        if not metin(r["kaynakca"]):
            uy.append(uyari("Orta", "Faktörün kaynağı yazılmamış", k, "Emisyon faktörünün kaynağını ve yılını (ör. ulusal envanter, tedarikçi beyanı) Kaynakça sütununa yazın."))
        fak[katla(k)] = {"kaynak": k, "birim": metin(r["birim"]), "faktor": f, "kapsam": int(kp), "kaynakca": metin(r["kaynakca"])}
    if not fak:
        raise ValueError(f"{fak_yol.name}: kullanılabilir emisyon faktörü yok.")
    for r in kayitlar(tuk_yol, TUK_SUTUN, ("kaynak", "miktar")):
        k = metin(r["kaynak"])
        if not k:
            continue
        kim, m, d = f"{k} · {metin(r['donem'])} (satır {r['_satir']})", para(r["miktar"]), donem(r["donem"])
        if m is None or m < 0:
            uy.append(uyari("Yüksek", "Tüketim miktarı okunamadı", kim, f"'{metin(r['miktar'])}'. Satır hesaba alınmadı."))
            continue
        if metin(r["donem"]) and not d:
            uy.append(uyari("Yüksek", "Dönem okunamadı", kim, f"'{metin(r['donem'])}'. AA.YYYY veya YYYY-AA yazın; satır hesaba alınmadı."))
            continue
        tuk.append({"donem": d, "tesis": metin(r["tesis"]) or "Belirtilmemiş", "kaynak": k, "miktar": m, "birim": metin(r["birim"]), "kim": kim})
    if ure_yol:
        for r in kayitlar(ure_yol, URE_SUTUN, ("miktar",)):
            m, d = para(r["miktar"]), donem(r["donem"])
            if m is None and not metin(r["donem"]):
                continue
            if m is None or m <= 0 or (metin(r["donem"]) and not d):
                uy.append(uyari("Yüksek", "Üretim satırı okunamadı", metin(r["donem"]), f"{ure_yol.name} satır {r['_satir']}: dönem '{metin(r['donem'])}', miktar '{metin(r['miktar'])}'."))
                continue
            ure[d] += m
    return tuk, fak, dict(ure), uy


def hesapla(tuk: list[dict], fak: dict, ure: dict | None = None) -> dict:
    uy, satirlar, eksik = [], [], defaultdict(int)
    for t in tuk:
        f = fak.get(katla(t["kaynak"]))
        if not f:
            eksik[t["kaynak"]] += 1
            continue
        if t["birim"] and f["birim"] and birim_katla(t["birim"]) != birim_katla(f["birim"]):
            uy.append(uyari("Yüksek", "Birim faktörle uyuşmuyor", t["kim"], f"Tüketim birimi '{t['birim']}', faktör birimi '{f['birim']}'. Birim çevrilmez; satır hesaba alınmadı."))
            continue
        satirlar.append(t | {"faktor": f["faktor"], "kapsam": f["kapsam"], "kaynak_ad": f["kaynak"], "emisyon": t["miktar"] * f["faktor"]})
    for k, n in eksik.items():
        uy.append(uyari("Yüksek", "Kaynağın emisyon faktörü yok", k, f"{n} tüketim satırı için faktör dosyasında karşılık yok; hesaba alınmadı."))
    if not satirlar:
        raise ValueError("Emisyon faktörü ve birimi eşleşen tüketim satırı yok.")
    ure = ure or {}
    top = sum((s["emisyon"] for s in satirlar), Z)

    def kapsamlar(l):
        return {k: sum((s["emisyon"] for s in l if s["kapsam"] == k), Z) for k in (1, 2, 3)}

    donemler = []
    for d in sorted({s["donem"] for s in satirlar}, key=lambda d: d or (0, 0)):
        l = [s for s in satirlar if s["donem"] == d]
        e, u = sum((s["emisyon"] for s in l), Z), ure.get(d)
        donemler.append({"donem": d, "kapsam": kapsamlar(l), "emisyon": e, "uretim": u, "urun_basina": yuvarla(e / u, "0.0001") if u else None})
    for d in ure:
        if d not in {x["donem"] for x in donemler}:
            uy.append(uyari("Orta", "Üretimi olan dönemin tüketimi yok", donem_ad(d), "Üretim dosyasında var, tüketim dosyasında bu döneme ait satır yok."))
    kaynak = {}
    for s in satirlar:
        x = kaynak.setdefault(katla(s["kaynak"]), {"kaynak": s["kaynak_ad"], "kapsam": s["kapsam"], "birim": s["birim"], "miktar": Z, "emisyon": Z, "faktor": s["faktor"]})
        x["miktar"] += s["miktar"]
        x["emisyon"] += s["emisyon"]
    kl = sorted(kaynak.values(), key=lambda x: -x["emisyon"])
    for x in kl:
        x["pay"] = yuvarla(x["emisyon"] / top * 100) if top else Z
    tesis = {}
    for s in satirlar:
        x = tesis.setdefault(s["tesis"], {"tesis": s["tesis"], "kapsam": {1: Z, 2: Z, 3: Z}, "emisyon": Z})
        x["kapsam"][s["kapsam"]] += s["emisyon"]
        x["emisyon"] += s["emisyon"]
    eslesen = [x for x in donemler if x["uretim"]]
    ut, et = sum((x["uretim"] for x in eslesen), Z), sum((x["emisyon"] for x in eslesen), Z)
    return {"satirlar": satirlar, "donemler": donemler, "kaynak": kl, "tesis": sorted(tesis.values(), key=lambda x: -x["emisyon"]), "kapsam": kapsamlar(satirlar), "toplam": top,
            "ton": yuvarla(top / 1000), "uretim": ut or None, "urun_basina": yuvarla(et / ut, "0.0001") if ut else None, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, fak: dict, uy: list[dict]) -> None:
    wb = Workbook()
    ton = lambda v: yuvarla(v / 1000, "0.001")   # noqa: E731
    ws = wb.active
    ws.title = "Özet"
    _baslik(ws, ["Gösterge", "Değer", "Birim"], (44, 18, 16))
    for ad, v, b in (("Kapsam 1 (tesiste yakılan yakıtlar)", ton(r["kapsam"][1]), "t CO2e"), ("Kapsam 2 (satın alınan elektrik / ısı)", ton(r["kapsam"][2]), "t CO2e"),
                     ("Kapsam 3 (girilen kalemler; kısmi)", ton(r["kapsam"][3]), "t CO2e"), ("Toplam", ton(r["toplam"]), "t CO2e"),
                     ("Üretim miktarı (tüketimi olan dönemler)", r["uretim"], "üretim birimi"), ("Ürün başına emisyon", r["urun_basina"], "kg CO2e / birim")):
        _satir(ws, [ad, v, b])
        ws.cell(ws.max_row, 2).number_format = "#,##0.000#"
    ws.append([])
    ws.append(["Emisyon faktörleri sizin verdiğiniz değerlerdir. Bu hesap doğrulanmış bir sera gazı envanteri değildir; kaçak ve proses emisyonları dahil değildir."])
    ws = wb.create_sheet("Dönem Bazında")
    _baslik(ws, ["Dönem", "Kapsam 1 (t CO2e)", "Kapsam 2 (t CO2e)", "Kapsam 3 (t CO2e)", "Toplam (t CO2e)", "Üretim", "Ürün Başına (kg CO2e)"], (12, 16, 16, 16, 16, 14, 18))
    for x in r["donemler"]:
        _satir(ws, [donem_ad(x["donem"]), ton(x["kapsam"][1]), ton(x["kapsam"][2]), ton(x["kapsam"][3]), ton(x["emisyon"]), x["uretim"], x["urun_basina"]])
        for j in (2, 3, 4, 5, 7):
            ws.cell(ws.max_row, j).number_format = "#,##0.000#"
    ws = wb.create_sheet("Kaynak Bazında")
    _baslik(ws, ["Kaynak", "Kapsam", "Tüketim", "Birim", "Emisyon Faktörü (kg CO2e / birim)", "Emisyon (t CO2e)", "Pay %"], (26, 9, 16, 9, 20, 16, 9))
    for x in r["kaynak"]:
        _satir(ws, [x["kaynak"], x["kapsam"], x["miktar"], x["birim"], x["faktor"], ton(x["emisyon"]), x["pay"]], para_s=(3, 7))
    ws = wb.create_sheet("Tesis Bazında")
    _baslik(ws, ["Tesis", "Kapsam 1 (t CO2e)", "Kapsam 2 (t CO2e)", "Kapsam 3 (t CO2e)", "Toplam (t CO2e)"], (26, 16, 16, 16, 16))
    for x in r["tesis"]:
        _satir(ws, [x["tesis"], ton(x["kapsam"][1]), ton(x["kapsam"][2]), ton(x["kapsam"][3]), ton(x["emisyon"])])
    ws = wb.create_sheet("Faktörler")
    _baslik(ws, ["Kaynak", "Birim", "Emisyon Faktörü (kg CO2e / birim)", "Kapsam", "Kaynakça"], (26, 9, 22, 9, 70))
    for f in fak.values():
        _satir(ws, [f["kaynak"], f["birim"], f["faktor"], f["kapsam"], f["kaynakca"]], dolgu={5: SARI if not f["kaynakca"] else None}, sar=(5,))
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(tuk_yol: Path, fak_yol: Path, cikti: Path, ure_yol: Path | None = None) -> dict:
    tuk, fak, ure, uy = oku(tuk_yol, fak_yol, ure_yol)
    r = hesapla(tuk, fak, ure)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, fak, r["uyarilar"])
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Tüketim verisi ve verdiğiniz emisyon faktörleriyle kapsam bazında emisyon ve ürün başına karbon değeri hesaplar.")
    p.add_argument("--tuketimler", type=Path, default=ORNEK / "tuketimler.csv", help="Dönem, Tesis, Kaynak, Miktar, Birim")
    p.add_argument("--faktorler", type=Path, default=ORNEK / "emisyon_faktorleri.csv", help="Kaynak, Birim, Emisyon Faktörü (kg CO2e / birim), Kapsam, Kaynakça")
    p.add_argument("--uretim", type=Path, help="Dönem, Tesis, Üretim Miktarı (ürün başına değer için)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "karbon_ayak_izi.xlsx")
    a = p.parse_args(argv)
    uretim = a.uretim or (ORNEK / "uretim.csv" if a.tuketimler == ORNEK / "tuketimler.csv" else None)
    for y in (a.tuketimler, a.faktorler, uretim):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.tuketimler, a.faktorler, a.cikti, uretim)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    k = r["kapsam"]
    print(f"[OK] Toplam {r['ton']} t CO2e · kapsam 1: {yuvarla(k[1] / 1000)} t · kapsam 2: {yuvarla(k[2] / 1000)} t" + (f" · kapsam 3: {yuvarla(k[3] / 1000)} t" if k[3] else "")
          + (f" · ürün başına {r['urun_basina']} kg CO2e" if r["urun_basina"] is not None else ""))
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['kim']} · {u['tur']}: {u['aciklama']}")
    print("[i] Emisyon faktörleri sizin verdiğiniz değerlerdir; kaynağını ve yılını doğrulayın.")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
