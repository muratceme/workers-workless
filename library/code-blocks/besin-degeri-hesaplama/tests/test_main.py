import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def b(ad, yag=0, doymus=0, karb=0, seker=0, lif=0, protein=0, tuz=0):
    return {"bilesen": ad, "yag": D(str(yag)), "doymus": D(str(doymus)), "karb": D(str(karb)), "seker": D(str(seker)), "lif": D(str(lif)), "protein": D(str(protein)), "tuz": D(str(tuz)),
            "poliol": D(0), "alkol": D(0), "asit": D(0)}


TABLO = {"un": b("Un", 1, 0.2, 72, 1.5, 3, 11), "seker": b("Şeker", karb=100, seker=100), "tereyagi": b("Tereyağı", 82, 52, 0.6, 0.6, 0, 0.7, 0.05), "yumurta": b("Yumurta", 10, 3, 0.7, 0.7, 0, 13, 0.35)}
RECETE = [{"bilesen": "Un", "miktar": D(500)}, {"bilesen": "Şeker", "miktar": D(100)}, {"bilesen": "Tereyağı", "miktar": D(200)}, {"bilesen": "yumurta", "miktar": D(100)}]


class HesapTesti(unittest.TestCase):
    def test_100g_ve_porsiyon(self):
        r = main.hesapla(RECETE, TABLO, D(800), porsiyon=D(50))
        self.assertEqual([(s["oge"], s["yuz_y"], s["porsiyon_y"], s["ra"]) for s in r["satirlar"]],
                         [("Enerji (kJ)", D(1972), D(986), D("11.7")), ("Enerji (kcal)", D(471), D(235), D("11.8")), ("Yağ", D("22.4"), D("11.2"), D("16.0")),
                          ("Doymuş yağ", D("13.5"), D("6.8"), D("33.8")), ("Karbonhidrat", D("57.7"), D("28.9"), D("11.1")), ("Şekerler", D("13.7"), D("6.8"), D("7.6")),
                          ("Lif", D("1.9"), D("0.9"), None), ("Protein", D("8.7"), D("4.3"), D("8.7")), ("Tuz", D("0.06"), D("0.03"), D("0.5"))])
        self.assertEqual((r["toplam_g"], r["agirlik"], r["katki"][0]["pay"], r["uyarilar"]), (D(900), D(800), D("55.56"), []))

    def test_kayip_ve_eksik(self):
        r = main.hesapla(RECETE + [{"bilesen": "Vanilya", "miktar": D(10)}], TABLO, kayip=D(20))
        self.assertEqual((r["agirlik"], r["satirlar"][2]["yuz_y"], [u["tur"] for u in r["uyarilar"]]), (D("728.0"), D("24.6"), ["Bileşen tabloda yok"]))
        r = main.hesapla(RECETE, TABLO, D(1000))
        self.assertEqual(([u["tur"] for u in r["uyarilar"]], r["satirlar"][0]["porsiyon_y"]), (["Son ağırlık reçeteden fazla"], None))
        for a, k in (([], {}), ([{"bilesen": "Un", "miktar": D(0)}], {})):
            with self.assertRaises(ValueError):
                main.hesapla(a, TABLO, **k)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "recete.csv", ORNEK / "besin_tablosu.csv", cikti, D(1000), porsiyon=D(40))
            self.assertEqual([(s["oge"], s["yuz_y"]) for s in r["satirlar"]],
                             [("Enerji (kJ)", D(1868)), ("Enerji (kcal)", D(446)), ("Yağ", D("20.2")), ("Doymuş yağ", D("12.2")), ("Karbonhidrat", D("55.6")), ("Şekerler", D("21.3")),
                              ("Lif", D("3.2")), ("Protein", D("8.8")), ("Tuz", D("0.57"))])
            self.assertEqual([(u["tur"], u["kim"]) for u in r["uyarilar"]], [])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Besin Değeri Tablosu", "Bileşen Katkısı", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, z = Path(tmp) / "r.csv", Path(tmp) / "t.csv"
            y.write_text("Hammadde;Gram\nUn;1.000\nSu;x\n;5\n", encoding="utf-8")
            z.write_text("Hammadde;Toplam Yağ;Doymuş;Karbonhidrat;Şekerler;Protein;Sodyum mg\nUn;1;2;70;80;10;400\nSu;;;;;;\nTuz;;;;;;x\nX;50;;60;;;\n", encoding="utf-8")
            l, uy = main.oku_recete(y)
            self.assertEqual(([(x["bilesen"], x["miktar"]) for x in l], len(uy)), ([("Un", D(1000))], 1))
            d, uy = main.oku_tablo(z)
            self.assertEqual(((d["un"]["tuz"], d["un"]["yag"], d["su"]["karb"]), [u["tur"] for u in uy]),
                             ((D("1.000"), D(1), D(0)), ["Tablo değeri tutarsız", "Tablo satırı okunamadı", "Tablo değeri tutarsız"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--recete", str(y), "--tablo", str(z), "--cikti", c]
            self.assertEqual(main.main(arg + ["--kayip", "10", "--porsiyon", "30", "--urun", "Deneme"]), 0)
            self.assertEqual(main.main(arg + ["--son-agirlik", "900"]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--son-agirlik", "0"], ["--kayip", "100"], ["--porsiyon", "x"])], [2] * 3)
            self.assertEqual(main.main(arg[:2] + ["--tablo", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Hammadde;Gram\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
