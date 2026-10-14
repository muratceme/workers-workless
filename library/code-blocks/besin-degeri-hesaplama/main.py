"""
Besin Değeri Hesaplama — Workers / Workless kod bloğu
Kalite Güvence ve Gıda Güvenliği › Gıda Mühendisi
Reçetedeki bileşenlerin besin değeri tablosundan 100 g (ve porsiyon) başına enerji ve besin öğelerini hesaplar.
  Besin öğesi / 100 g = Σ(bileşen miktarı × bileşenin 100 g'daki değeri ÷ 100) ÷ son ürün ağırlığı × 100
  Son ürün ağırlığı   = --son-agirlik; verilmezse reçete toplamı × (1 − --kayip %)   (pişirmede su kaybı)
  Enerji              = yağ × 37 kJ (9 kcal) + karbonhidrat × 17 (4) + protein × 17 (4) + lif × 8 (2) + poliol × 10 (2,4)
                        + alkol × 29 (7) + organik asit × 13 (3)   — g başına; karbonhidrat lif ve poliol hariç
  Tuz                 = tuz sütunu; yoksa sodyum × 2,5
  Referans alım %     = porsiyondaki miktar ÷ referans alım × 100 (enerji 8400 kJ / 2000 kcal, yağ 70 g, doymuş yağ 20 g,
                        karbonhidrat 260 g, şeker 90 g, protein 50 g, tuz 6 g)
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

RECETE = {"bilesen": ("bilesen", "hammadde", "malzeme", "icerik"), "miktar": ("miktar", "miktar g", "gram", "agirlik")}
TABLO = {"bilesen": ("bilesen", "hammadde", "malzeme", "icerik"), "yag": ("yag", "toplam yag"), "doymus": ("doymus yag", "doymus yag asitleri", "doymus"),
         "karb": ("karbonhidrat", "karbonhidratlar"), "seker": ("seker", "sekerler"), "lif": ("lif", "diyet lifi", "posa"), "protein": ("protein",), "tuz": ("tuz",),
         "sodyum": ("sodyum", "sodyum g"), "sodyum_mg": ("sodyum mg",), "poliol": ("poliol", "polioller"), "alkol": ("alkol", "etanol"), "asit": ("organik asit", "organik asitler")}
OGE = ("yag", "doymus", "karb", "seker", "lif", "protein", "tuz", "poliol", "alkol", "asit")
AD = {"yag": "Yağ", "doymus": "Doymuş yağ", "karb": "Karbonhidrat", "seker": "Şekerler", "lif": "Lif", "protein": "Protein", "tuz": "Tuz", "poliol": "Polioller", "alkol": "Alkol",
      "asit": "Organik asitler"}
KJ = {"yag": 37, "karb": 17, "protein": 17, "lif": 8, "poliol": 10, "alkol": 29, "asit": 13}
KCAL = {"yag": Decimal(9), "karb": Decimal(4), "protein": Decimal(4), "lif": Decimal(2), "poliol": Decimal("2.4"), "alkol": Decimal(7), "asit": Decimal(3)}
RA = {"kj": Decimal(8400), "kcal": Decimal(2000), "yag": Decimal(70), "doymus": Decimal(20), "karb": Decimal(260), "seker": Decimal(90), "protein": Decimal(50), "tuz": Decimal(6)}
SIFIR = Decimal(0)


def oku_recete(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, RECETE, ("bilesen", "miktar")):
        ad = metin(r["bilesen"])
        if not ad:
            continue
        m = para(r["miktar"])
        if m is None or m < 0:
            uy.append(uyari("Yüksek", "Reçete satırı okunamadı", ad, f"Satır {r['_satir']}: miktar (g) okunamadı. Bileşen alınmadı."))
            continue
        l.append({"bilesen": ad, "miktar": m})
    return l, uy


def oku_tablo(yol: Path) -> tuple[dict, list[dict]]:
    uy, d = [], {}
    for r in kayitlar(yol, TABLO, ("bilesen",)):
        ad = metin(r["bilesen"])
        if not ad:
            continue
        v = {a: para(r[a]) if metin(r[a]) else None for a in OGE + ("sodyum", "sodyum_mg")}
        if any(metin(r[a]) and (v[a] is None or v[a] < 0) for a in v):
            uy.append(uyari("Yüksek", "Tablo satırı okunamadı", ad, f"Satır {r['_satir']}: besin değerlerinden biri okunamadı. Bileşen alınmadı."))
            continue
        if v["tuz"] is None:
            na = v["sodyum"] if v["sodyum"] is not None else v["sodyum_mg"] / 1000 if v["sodyum_mg"] is not None else None
            v["tuz"] = na * Decimal("2.5") if na is not None else None
        deger = {a: v[a] or SIFIR for a in OGE}
        ad_k = katla(ad)
        if deger["doymus"] > deger["yag"] or deger["seker"] > deger["karb"]:
            uy.append(uyari("Orta", "Tablo değeri tutarsız", ad, "Doymuş yağ toplam yağdan ya da şeker karbonhidrattan büyük."))
        if deger["yag"] + deger["karb"] + deger["protein"] + deger["lif"] + deger["tuz"] > 100:
            uy.append(uyari("Orta", "Tablo değeri tutarsız", ad, "Besin öğelerinin toplamı 100 g'ı aşıyor; birimleri kontrol edin."))
        d[ad_k] = {"bilesen": ad, **deger}
    return d, uy


def hesapla(recete: list[dict], tablo: dict, son_agirlik: Decimal | None = None, kayip: Decimal = SIFIR, porsiyon: Decimal | None = None) -> dict:
    if not recete:
        raise ValueError("Reçetede bileşen yok.")
    uy, katki = [], []
    toplam_g = sum((b["miktar"] for b in recete), SIFIR)
    if toplam_g <= 0:
        raise ValueError("Reçete toplamı sıfır.")
    agirlik = son_agirlik if son_agirlik else toplam_g * (1 - kayip / 100)
    if agirlik <= 0:
        raise ValueError("Son ürün ağırlığı sıfırdan büyük olmalı.")
    toplam = {a: SIFIR for a in OGE}
    for b in recete:
        t = tablo.get(katla(b["bilesen"]))
        if t is None:
            uy.append(uyari("Yüksek", "Bileşen tabloda yok", b["bilesen"], f"{b['miktar']:,.1f} g; besin değeri bilinmediği için hesaba sıfır katıldı. Sonuç eksik."))
            katki.append({"bilesen": b["bilesen"], "miktar": b["miktar"], "pay": yuvarla(b["miktar"] * 100 / toplam_g), **{a: None for a in OGE}})
            continue
        k = {a: t[a] * b["miktar"] / 100 for a in OGE}
        for a in OGE:
            toplam[a] += k[a]
        katki.append({"bilesen": b["bilesen"], "miktar": b["miktar"], "pay": yuvarla(b["miktar"] * 100 / toplam_g), **k})
    if son_agirlik and son_agirlik > toplam_g:
        uy.append(uyari("Orta", "Son ağırlık reçeteden fazla", "", f"Son ürün {son_agirlik:,.1f} g, reçete toplamı {toplam_g:,.1f} g. Su ekleniyorsa reçeteye yazın."))
    yuz = {a: toplam[a] * 100 / agirlik for a in OGE}
    kj = sum((yuz[a] * KJ[a] for a in KJ), SIFIR)
    kcal = sum((yuz[a] * KCAL[a] for a in KCAL), SIFIR)
    satirlar = [{"oge": "Enerji (kJ)", "anahtar": "kj", "yuz": kj}, {"oge": "Enerji (kcal)", "anahtar": "kcal", "yuz": kcal}] + \
               [{"oge": AD[a], "anahtar": a, "yuz": yuz[a]} for a in OGE if a not in ("poliol", "alkol", "asit") or toplam[a]]
    for s in satirlar:
        s["porsiyon"] = s["yuz"] * porsiyon / 100 if porsiyon else None
        s["ra"] = yuvarla(s["porsiyon"] * 100 / RA[s["anahtar"]], "0.1") if porsiyon and s["anahtar"] in RA else None
        hane = "1" if s["anahtar"] in ("kj", "kcal") else "0.01" if s["anahtar"] == "tuz" else "0.1"
        s["yuz_y"] = yuvarla(s["yuz"], hane)
        s["porsiyon_y"] = yuvarla(s["porsiyon"], hane) if porsiyon else None
    return {"satirlar": satirlar, "katki": katki, "toplam_g": toplam_g, "agirlik": agirlik, "porsiyon": porsiyon, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, urun: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Besin Değeri Tablosu"
    p = r["porsiyon"]
    _baslik(ws, ["Besin öğesi", "100 g'da"] + ([f"Porsiyon ({p:,.0f} g)", "Porsiyonda referans alım %"] if p else []), (22, 12, 16, 16))
    for s in r["satirlar"]:
        _satir(ws, [s["oge"], s["yuz_y"]] + ([s["porsiyon_y"], s["ra"]] if p else []))
    ws.append([])
    ws.append([f"{urun or 'Ürün'}: reçete {r['toplam_g']:,.1f} g → son ürün {r['agirlik']:,.1f} g. Hesaplanmış değerdir; etiket için laboratuvar analizi ve yuvarlama kurallarıyla doğrulayın."])
    ws.append(["Referans alım: enerji 8400 kJ / 2000 kcal, yağ 70 g, doymuş yağ 20 g, karbonhidrat 260 g, şekerler 90 g, protein 50 g, tuz 6 g (ortalama yetişkin)."])
    ws = wb.create_sheet("Bileşen Katkısı")
    _baslik(ws, ["Bileşen", "Miktar (g)", "Reçetedeki Pay %"] + [f"{AD[a]} (g)" for a in OGE], (24, 10, 10) + (10,) * len(OGE), "B2")
    for k in r["katki"]:
        _satir(ws, [k["bilesen"], k["miktar"], k["pay"]] + [yuvarla(k[a], "0.01") if k[a] is not None else None for a in OGE], para_s=range(3, 4 + len(OGE)),
               dolgu={1: KIRMIZI if k["yag"] is None else None})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(recete_yol: Path, tablo_yol: Path, cikti: Path, son_agirlik: Decimal | None = None, kayip: Decimal = SIFIR, porsiyon: Decimal | None = None, urun: str = "") -> dict:
    recete, uy1 = oku_recete(recete_yol)
    tablo, uy2 = oku_tablo(tablo_yol)
    r = hesapla(recete, tablo, son_agirlik, kayip, porsiyon)
    r["uyarilar"] = sirala(uy1 + uy2 + r["uyarilar"])
    rapor_yaz(cikti, r, urun)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Reçeteden 100 g ve porsiyon başına enerji ve besin öğelerini hesaplar.")
    p.add_argument("--recete", type=Path, help="Bileşen, Miktar (g)")
    p.add_argument("--tablo", type=Path, help="Bileşen, Yağ, Doymuş Yağ, Karbonhidrat, Şeker, Lif, Protein, Tuz (ya da Sodyum) — 100 g'daki değerler")
    p.add_argument("--son-agirlik", help="Pişmiş / işlenmiş son ürün ağırlığı, g")
    p.add_argument("--kayip", default="0", help="Son ağırlık verilmezse işlemdeki ağırlık kaybı %% (varsayılan 0)")
    p.add_argument("--porsiyon", help="Porsiyon, g (verilirse porsiyon ve referans alım %% sütunları eklenir)")
    p.add_argument("--urun", default="", help="Ürün adı (rapor notunda)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "besin_degeri.xlsx")
    a = p.parse_args(argv)
    ornek = a.recete is None
    ry, ty = a.recete or ORNEK / "recete.csv", a.tablo or ORNEK / "besin_tablosu.csv"
    if ornek and not (a.son_agirlik or a.porsiyon):
        a.son_agirlik, a.porsiyon, a.urun = "1.000", "40", a.urun or "Kakaolu kurabiye"
    son = para(a.son_agirlik) if a.son_agirlik else None
    kayip, pors = para(a.kayip), (para(a.porsiyon) if a.porsiyon else None)
    if (a.son_agirlik and (son is None or son <= 0)) or kayip is None or not 0 <= kayip < 100 or (a.porsiyon and (pors is None or pors <= 0)):
        print("[X] --son-agirlik ve --porsiyon sıfırdan büyük, --kayip 0-100 arası olmalı")
        return 2
    for y in (ry, ty):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(ry, ty, a.cikti, son, kayip, pors, a.urun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(r['katki'])} bileşen · reçete {r['toplam_g']:,.1f} g → son ürün {r['agirlik']:,.1f} g")
    for s in r["satirlar"]:
        print(f"     {s['oge']}: {s['yuz_y']} / 100 g" + (f" · porsiyon {s['porsiyon_y']}" + (f" (RA %{s['ra']})" if s["ra"] is not None else "") if s["porsiyon_y"] is not None else ""))
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
