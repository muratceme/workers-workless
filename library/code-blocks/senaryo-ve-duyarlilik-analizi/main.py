"""
Senaryo ve Duyarlılık Analizi — Workers / Workless kod bloğu
Finans › Finansal Analist

Kur, faiz, satış ve maliyet varsayımlarından yıllık kârlılık ve nakit senaryoları üretir:
  - Model (sürücü tabanlı):
      gelir          = yurt içi hacim × yurt içi fiyat + ihracat hacmi × ihracat fiyatı (EUR) × EUR/TRY
      değişken maliyet = hacim × (ithal hammadde (USD) × USD/TRY + yerli değişken maliyet)
      FAVÖK          = gelir − değişken maliyet − personel − diğer sabit giderler
      faiz gideri    = TL kredi × TL faiz + EUR kredi × EUR faiz × EUR/TRY
      kur farkı      = net EUR pozisyonu × (EUR/TRY − dönem başı EUR/TRY)
      vergi öncesi kâr = FAVÖK − amortisman − faiz + kur farkı;  vergi = max(0, VÖK) × vergi oranı
      serbest nakit (basit) = FAVÖK − faiz − vergi − yatırım harcaması (+ kur farkı nakit değildir)
  - Senaryolar: parametre bazında "%" (göreli), "puan" (oranlara eklenir) veya "=" (değer atanır).
  - Duyarlılık (tornado): her parametre tek tek tutarlarda ±%10, oranlarda ±2 puan değiştirilir; net kâra etkisi
    büyükten küçüğe sıralanır.
  - Başa baş: net kârı sıfırlayan satış hacmi ve EUR/TRY kuru (diğerleri sabit; ikiye bölme).
Rapor: senaryo karşılaştırması, tornado, başa baş, varsayımlar, uyarılar. İnternete bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: ihracatçı üretici, 4 senaryo
    python main.py --varsayimlar v.csv --senaryolar s.csv --adim 10 --oran-adim 2
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import OrderedDict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")

# anahtar → (görünen ad, tür) · tür: "oran" (0–1) veya "tutar"
PARAMETRELER = OrderedDict([
    ("hacim", ("Satış hacmi", "tutar")), ("ihracat_pay", ("İhracat payı", "oran")), ("ic_fiyat", ("Yurt içi birim fiyat", "tutar")),
    ("ihr_fiyat", ("İhracat birim fiyatı", "tutar")), ("eur", ("EUR/TRY", "tutar")), ("eur_bas", ("EUR/TRY dönem başı", "tutar")),
    ("usd", ("USD/TRY", "tutar")), ("hammadde_usd", ("Birim ithal hammadde", "tutar")), ("yerli_maliyet", ("Birim yerli değişken maliyet", "tutar")),
    ("personel", ("Personel gideri", "tutar")), ("sabit", ("Diğer sabit giderler", "tutar")), ("amortisman", ("Amortisman", "tutar")),
    ("tl_kredi", ("TL kredi", "tutar")), ("tl_faiz", ("TL faiz oranı", "oran")), ("eur_kredi", ("EUR kredi", "tutar")),
    ("eur_faiz", ("EUR faiz oranı", "oran")), ("eur_poz", ("Net EUR pozisyonu", "tutar")), ("yatirim", ("Yatırım harcaması", "tutar")),
    ("vergi", ("Vergi oranı", "oran")),
])
AD_ANAHTAR = {}
CIKTILAR = [("gelir", "Gelir"), ("degisken", "Değişken maliyet"), ("brut", "Brüt katkı"), ("favok", "FAVÖK"), ("amortisman", "Amortisman"),
            ("faiz", "Faiz gideri"), ("kur_farki", "Kur farkı (net pozisyon)"), ("vok", "Vergi öncesi kâr"), ("vergi_t", "Vergi"), ("net", "Net kâr"),
            ("nakit", "Serbest nakit (basit)"), ("favok_marj", "FAVÖK marjı"), ("net_marj", "Net kâr marjı"), ("faiz_karsilama", "Faiz karşılama (FAVÖK ÷ faiz)")]


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


for _k, (_ad, _) in PARAMETRELER.items():
    AD_ANAHTAR[katla(_ad)] = _k


def sayi(x) -> float | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip().replace("TL", "").replace("%", "").replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return float(Decimal(s))
    except InvalidOperation:
        return None


def tl(x: float) -> str:
    return f"{x:,.0f}".replace(",", ".")


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


# ----------------------------------------------------------------------------
# Model
# ----------------------------------------------------------------------------

def model(p: dict) -> dict:
    g = lambda k: p.get(k, 0.0)  # noqa: E731
    dis = g("hacim") * g("ihracat_pay")
    ic = g("hacim") - dis
    o = {"gelir": ic * g("ic_fiyat") + dis * g("ihr_fiyat") * g("eur"), "degisken": g("hacim") * (g("hammadde_usd") * g("usd") + g("yerli_maliyet"))}
    o["brut"] = o["gelir"] - o["degisken"]
    o["favok"] = o["brut"] - g("personel") - g("sabit")
    o["amortisman"] = g("amortisman")
    o["faiz"] = g("tl_kredi") * g("tl_faiz") + g("eur_kredi") * g("eur_faiz") * g("eur")
    o["kur_farki"] = g("eur_poz") * (g("eur") - (g("eur_bas") or g("eur")))
    o["vok"] = o["favok"] - o["amortisman"] - o["faiz"] + o["kur_farki"]
    o["vergi_t"] = max(o["vok"], 0) * g("vergi")
    o["net"] = o["vok"] - o["vergi_t"]
    o["nakit"] = o["favok"] - o["faiz"] - o["vergi_t"] - g("yatirim")
    o["favok_marj"] = o["favok"] / o["gelir"] if o["gelir"] else None
    o["net_marj"] = o["net"] / o["gelir"] if o["gelir"] else None
    o["faiz_karsilama"] = o["favok"] / o["faiz"] if o["faiz"] else None
    return o


def oku_varsayimlar(yol: Path) -> tuple[dict, list[dict]]:
    p, uy = {}, []
    for r in tablo_oku(yol)[1:]:
        if len(r) < 2 or not r[0]:
            continue
        k = AD_ANAHTAR.get(katla(r[0]))
        v = sayi(r[1])
        if k is None:
            uy.append({"onem": "Orta", "tur": "Tanınmayan parametre", "aciklama": f"'{r[0]}' modelde yok; kullanılmadı"})
            continue
        if v is None:
            uy.append({"onem": "Orta", "tur": "Okunamayan değer", "aciklama": f"'{r[0]}': '{r[1]}'"})
            continue
        if PARAMETRELER[k][1] == "oran" and (abs(v) > 1 or "%" in str(r[1])):
            v /= 100
        p[k] = v
    for k, (ad, _) in PARAMETRELER.items():
        if k not in p and k not in ("eur_bas", "eur_poz", "eur_kredi", "eur_faiz", "tl_kredi", "tl_faiz", "yatirim", "amortisman", "ihracat_pay"):
            uy.append({"onem": "Orta", "tur": "Eksik parametre", "aciklama": f"'{ad}' verilmedi; 0 kabul edildi"})
    return p, uy


def uygula(p: dict, degisiklikler: list[tuple[str, str, float]]) -> dict:
    q = dict(p)
    for k, tip, v in degisiklikler:
        oran = PARAMETRELER[k][1] == "oran"
        if tip == "%":
            q[k] = q.get(k, 0.0) * (1 + v / 100)
        elif tip == "puan":
            q[k] = q.get(k, 0.0) + (v / 100 if oran else v)
        else:
            q[k] = v / 100 if oran and abs(v) > 1 else v
    return q


def oku_senaryolar(yol: Path | None) -> tuple[OrderedDict, list[dict]]:
    sen, uy = OrderedDict(), []
    if not yol:
        return sen, uy
    s = tablo_oku(yol)
    for r in s[1:]:
        if len(r) < 4 or not r[0]:
            continue
        k = AD_ANAHTAR.get(katla(r[1]))
        tip = "%" if "%" in str(r[2]) or katla(r[2]).startswith(("yuzde", "oran")) else "puan" if katla(r[2]).startswith("puan") else "=" if "=" in str(r[2]) \
            or katla(r[2]).startswith(("deger", "ata")) else None
        v = sayi(r[3])
        if k is None or tip is None or v is None:
            uy.append({"onem": "Orta", "tur": "Senaryo satırı", "aciklama": f"'{r[0]}' · '{r[1]}' · '{r[2]}' · '{r[3]}' anlaşılamadı; atlandı"})
            continue
        sen.setdefault(str(r[0]).strip(), []).append((k, tip, v))
    return sen, uy


def basa_bas(p: dict, k: str, alt: float, ust: float) -> float | None:
    f = lambda x: model({**p, k: x})["net"]  # noqa: E731
    fa, fu = f(alt), f(ust)
    if fa * fu > 0:
        return None
    for _ in range(100):
        m = (alt + ust) / 2
        fm = f(m)
        if fa * fm <= 0:
            ust = m
        else:
            alt, fa = m, fm
    return (alt + ust) / 2


def analiz_et(p: dict, senaryolar: OrderedDict, adim: float, oran_adim: float) -> dict:
    uy = []
    temel = model(p)
    sonuc = OrderedDict([("Temel", temel)])
    for ad, d in senaryolar.items():
        sonuc[ad] = model(uygula(p, d))
    tornado = []
    for k, (ad, tur) in PARAMETRELER.items():
        if k not in p or (tur == "tutar" and p[k] == 0):
            continue
        if tur == "oran":
            dus, art = uygula(p, [(k, "puan", -oran_adim)]), uygula(p, [(k, "puan", oran_adim)])
            etiket = f"±{oran_adim:g} puan"
        else:
            dus, art = uygula(p, [(k, "%", -adim)]), uygula(p, [(k, "%", adim)])
            etiket = f"±%{adim:g}"
        n1, n2 = model(dus)["net"], model(art)["net"]
        tornado.append({"k": k, "ad": ad, "etiket": etiket, "dusuk": n1 - temel["net"], "yuksek": n2 - temel["net"], "aralik": abs(n2 - n1)})
    tornado.sort(key=lambda x: -x["aralik"])
    bb = {"hacim": basa_bas(p, "hacim", 0, p.get("hacim", 0) * 10) if p.get("hacim") else None,
          "eur": basa_bas(p, "eur", p.get("eur", 0) * 0.2, p.get("eur", 0) * 5) if p.get("eur") else None}
    if temel["net"] < 0:
        uy.append({"onem": "Yüksek", "tur": "Temel senaryo zarar", "aciklama": f"Temel senaryoda net zarar {tl(temel['net'])} TL"})
    for ad, o in sonuc.items():
        if o["faiz_karsilama"] is not None and o["faiz_karsilama"] < 1.5:
            uy.append({"onem": "Orta", "tur": "Faiz karşılama", "aciklama": f"{ad}: FAVÖK faizi {o['faiz_karsilama']:.2f} kez karşılıyor (< 1,5)".replace(".", ",", 1)})
        if o["nakit"] < 0:
            uy.append({"onem": "Orta", "tur": "Negatif nakit", "aciklama": f"{ad}: serbest nakit {tl(o['nakit'])} TL; finansman ihtiyacı doğar"})
    if bb["hacim"] is not None and p.get("hacim") and bb["hacim"] > p["hacim"] * 0.9:
        uy.append({"onem": "Orta", "tur": "Düşük güvenlik payı", "aciklama": f"Başa baş hacim {tl(bb['hacim'])}; mevcut hacmin %{bb['hacim'] / p['hacim'] * 100:.0f}'i"})
    if p.get("eur_poz", 0) < 0:
        uy.append({"onem": "Bilgi", "tur": "Kısa döviz pozisyonu", "aciklama": f"Net EUR pozisyonu {tl(p['eur_poz'])} EUR; kur artışı kur farkı zararı yazar "
                   "(ihracat gelirindeki artış ayrıca gelire yansır)"})
    return {"p": p, "sonuc": sonuc, "tornado": tornado, "basa_bas": bb, "uyarilar": uy, "senaryolar": senaryolar}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
UST = Alignment(vertical="top", wrap_text=True)
PF = "#,##0;[Red]-#,##0"


def _baslik(ws, basliklar, genislik=()):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def rapor_yaz(cikti: Path, s: dict) -> None:
    wb = Workbook()
    sk = wb.active
    sk.title = "Senaryolar"
    adlar = list(s["sonuc"])
    _baslik(sk, ["Kalem"] + adlar + [f"{a} − Temel" for a in adlar[1:]], [30] + [16] * (2 * len(adlar) - 1))
    temel = s["sonuc"]["Temel"]
    for k, ad in CIKTILAR:
        satir = [ad] + [s["sonuc"][a][k] for a in adlar]
        if not k.endswith(("marj", "karsilama")):
            satir += [s["sonuc"][a][k] - temel[k] for a in adlar[1:]]
        sk.append(satir)
        fmt = "0.0%" if k.endswith("marj") else "0.00" if k == "faiz_karsilama" else PF
        for j in range(2, sk.max_column + 1):
            sk.cell(sk.max_row, j).number_format = fmt
        if k in ("net", "favok", "nakit"):
            for h in sk[sk.max_row]:
                h.font = Font(bold=True)
    sk.freeze_panes = "B2"
    g = BarChart()
    g.title, g.height, g.width = "Net kâr ve serbest nakit", 8, 18
    ni = [k for k, _ in CIKTILAR].index("net") + 2
    g.add_data(Reference(sk, min_col=1, max_col=1 + len(adlar), min_row=ni, max_row=ni), titles_from_data=True, from_rows=True)
    g.add_data(Reference(sk, min_col=1, max_col=1 + len(adlar), min_row=ni + 1, max_row=ni + 1), titles_from_data=True, from_rows=True)
    g.set_categories(Reference(sk, min_col=2, max_col=1 + len(adlar), min_row=1))
    sk.add_chart(g, f"A{sk.max_row + 3}")
    sk.append([])
    sk.append(["Senaryo tanımları"])
    for ad, d in s["senaryolar"].items():
        sk.append([ad, "; ".join(f"{PARAMETRELER[k][0]} {('+' if v > 0 else '') + format(v, 'g')}{'%' if t == '%' else ' puan' if t == 'puan' else ''}"
                                 if t != "=" else f"{PARAMETRELER[k][0]} = {v:g}" for k, t, v in d)])

    tr = wb.create_sheet("Tornado")
    _baslik(tr, ["Parametre", "Değişim", "Düşüşte Net Kâr Etkisi", "Artışta Net Kâr Etkisi", "Aralık"], (30, 12, 20, 20, 16))
    for x in s["tornado"]:
        tr.append([x["ad"], x["etiket"], x["dusuk"], x["yuksek"], x["aralik"]])
        for j in (3, 4, 5):
            tr.cell(tr.max_row, j).number_format = PF
    if s["tornado"]:
        g = BarChart()
        g.type, g.grouping, g.overlap = "bar", "clustered", 100
        g.title, g.height, g.width = "Net kâra etki (tornado)", 10, 18
        g.add_data(Reference(tr, min_col=3, max_col=4, min_row=1, max_row=min(tr.max_row, 11)), titles_from_data=True)
        g.set_categories(Reference(tr, min_col=1, min_row=2, max_row=min(tr.max_row, 11)))
        g.y_axis.scaling.orientation = "maxMin"
        tr.add_chart(g, "G2")

    bb = wb.create_sheet("Başa Baş")
    _baslik(bb, ["Değişken", "Mevcut", "Başa Baş Değeri", "Güvenlik Payı", "Açıklama"], (20, 14, 16, 14, 60))
    p = s["p"]
    for k, ad, aciklama in (("hacim", "Satış hacmi", "Net kârı sıfırlayan hacim (diğer varsayımlar sabit)"),
                            ("eur", "EUR/TRY", "Net kârı sıfırlayan EUR/TRY (USD/TRY sabit; ihracat geliri, EUR faiz ve kur farkı birlikte)")):
        v = s["basa_bas"][k]
        bb.append([ad, p.get(k), v if v is not None else "Bulunamadı", (p[k] - v) / p[k] if v is not None and p.get(k) else None, aciklama])
        bb.cell(bb.max_row, 2).number_format = bb.cell(bb.max_row, 3).number_format = "#,##0.00"
        bb.cell(bb.max_row, 4).number_format = "0.0%"

    vs = wb.create_sheet("Varsayımlar")
    _baslik(vs, ["Parametre", "Temel Değer"] + adlar[1:], [30, 16] + [16] * (len(adlar) - 1))
    senaryo_p = {a: uygula(p, d) for a, d in s["senaryolar"].items()}
    for k, (ad, tur) in PARAMETRELER.items():
        vs.append([ad, p.get(k)] + [senaryo_p[a].get(k) for a in adlar[1:]])
        for j in range(2, vs.max_column + 1):
            c = vs.cell(vs.max_row, j)
            c.number_format = "0.00%" if tur == "oran" else "#,##0.00"
            if j > 2 and c.value != p.get(k):
                c.fill = PatternFill("solid", fgColor="FFF4CE")

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Açıklama"], (9, 24, 100))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["aciklama"]])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 3).alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(varsayim_yolu: Path, cikti: Path, senaryo_yolu: Path | None = None, adim: float = 10, oran_adim: float = 2) -> dict:
    p, u1 = oku_varsayimlar(varsayim_yolu)
    if not p:
        raise ValueError(f"{varsayim_yolu.name}: parametre bulunamadı")
    sen, u2 = oku_senaryolar(senaryo_yolu)
    s = analiz_et(p, sen, adim, oran_adim)
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    s["uyarilar"] = sorted(u1 + u2 + s["uyarilar"], key=lambda u: sira[u["onem"]])
    rapor_yaz(cikti, s)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Kur, faiz, satış ve maliyet varsayımlarından kârlılık ve nakit senaryoları, tornado ve başa baş analizi üretir.")
    p.add_argument("--varsayimlar", type=Path, default=ORNEK / "varsayimlar.csv", help="Parametre;Değer listesi (bkz. README)")
    p.add_argument("--senaryolar", type=Path, help="Senaryo;Parametre;Değişim (%% / puan / =);Değer")
    p.add_argument("--adim", type=float, default=10, help="Tornado: tutarlarda ±%% değişim (varsayılan 10)")
    p.add_argument("--oran-adim", type=float, default=2, help="Tornado: oranlarda ± puan değişim (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "senaryo_analizi.xlsx")
    a = p.parse_args(argv)
    sen = a.senaryolar or (ORNEK / "senaryolar.csv" if a.varsayimlar == ORNEK / "varsayimlar.csv" else None)
    for y in (a.varsayimlar, sen):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        s = calistir(a.varsayimlar, a.cikti, sen, a.adim, a.oran_adim)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    for ad, o in s["sonuc"].items():
        print(f"[OK] {ad}: gelir {tl(o['gelir'])} · FAVÖK {tl(o['favok'])} · net kâr {tl(o['net'])} · serbest nakit {tl(o['nakit'])} TL")
    if s["tornado"]:
        print("     En etkili varsayımlar: " + ", ".join(f"{x['ad']} ({tl(x['aralik'])})" for x in s["tornado"][:3]))
    for u in s["uyarilar"]:
        if u["onem"] != "Bilgi":
            print(f"[!] {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
