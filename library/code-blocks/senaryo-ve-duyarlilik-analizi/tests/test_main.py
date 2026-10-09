import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "s.xlsx"
        cls.s = main.calistir(ORNEK / "varsayimlar.csv", cls.cikti, ORNEK / "senaryolar.csv")

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_temel_model(self):                          # elle: 72.000 × 2.350 + 48.000 × 52 × 48,5 = 290.256.000
        t = self.s["sonuc"]["Temel"]
        self.assertAlmostEqual(t["gelir"], 290256000)
        self.assertAlmostEqual(t["degisken"], 152688000)   # 120.000 × (18 × 41,8 + 520)
        self.assertAlmostEqual(t["faiz"], 21892500)        # 40M × %42 + 1,5M × %7 × 48,5
        self.assertAlmostEqual(t["kur_farki"], -1350000)   # −900.000 × (48,5 − 47)
        self.assertAlmostEqual(t["net"], 34369125)

    def test_senaryolar(self):
        k = self.s["sonuc"]["Kur şoku"]
        self.assertAlmostEqual(k["kur_farki"], -900000 * (48.5 * 1.25 - 47))
        self.assertLess(self.s["sonuc"]["Kötümser"]["net"], self.s["sonuc"]["Temel"]["net"])
        p = main.uygula(self.s["p"], self.s["senaryolar"]["İhracat odaklı"])
        self.assertAlmostEqual(p["ihracat_pay"], 0.55)
        p = main.uygula(self.s["p"], self.s["senaryolar"]["Kötümser"])
        self.assertAlmostEqual(p["tl_faiz"], 0.50)

    def test_tornado_ve_basa_bas(self):
        t = self.s["tornado"][0]
        self.assertEqual(t["ad"], "Yurt içi birim fiyat")
        self.assertAlmostEqual(t["yuksek"], 72000 * 235 * 0.75)
        bb = self.s["basa_bas"]["hacim"]
        self.assertAlmostEqual(main.model({**self.s["p"], "hacim": bb})["net"], 0, delta=1)
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Senaryolar", "Tornado", "Başa Baş", "Varsayımlar", "Uyarılar"])


class KuralTesti(unittest.TestCase):
    def test_hatali_girdiler(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t / "v.csv").write_text("Parametre;Değer\nSatış hacmi;1000\nYurt içi birim fiyat;100\nUçuş hızı;5\nVergi oranı;25\n", encoding="utf-8")
            (t / "s.csv").write_text("Senaryo;Parametre;Değişim;Değer\nA;Satış hacmi;%;-50\nB;Bilinmeyen;%;5\n", encoding="utf-8")
            s = main.calistir(t / "v.csv", t / "o.xlsx", t / "s.csv")
            turler = {u["tur"] for u in s["uyarilar"]}
            self.assertTrue({"Tanınmayan parametre", "Senaryo satırı", "Eksik parametre"} <= turler)
            self.assertAlmostEqual(s["p"]["vergi"], 0.25)
            self.assertAlmostEqual(s["sonuc"]["A"]["gelir"], 50000)

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--varsayimlar", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
