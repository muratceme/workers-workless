"""
Hakediş Hesaplama — Workers / Workless kod bloğu
İnşaat › Teknik Ofis Mühendisi

Birim fiyatlı sözleşmenin poz listesi ve dönem metrajlarından hakediş raporu hazırlar: kümülatif metraj
(yeşil defter özeti), hakediş icmali (toplam − önceki = bu hakediş), KDV, KDV tevkifatı, gelir/kurumlar
vergisi stopajı, teminat ve avans mahsubu kesintileri ve ödenecek net tutar. Sözleşme miktarını aşan pozları
ve toplam iş artışını işaretler. İnternete bağlanmaz.

Varsayılan vergi kuralları (sozlesme.json ile değiştirilebilir, ödemeden önce mali müşavirinizle teyit edin):
  - KDV tevkifatı 4/10: belirlenmiş alıcılara ya da KDV dahil bedeli 5.000.000 TL ve üzeri yapım işlerinde
    (KDV Genel Uygulama Tebliği I/C-2.1.3.2.1; Seri No: 35 ile 01.03.2021'den itibaren)
  - Stopaj %5: yıllara yaygın inşaat ve onarım işleri hakedişlerinde (GVK md. 94/3, KVK md. 15; 3491 sayılı
    CBK ile 01.03.2021'den itibaren). Demiryolu, gemi ve nükleer santral inşaatında oran %1'dir.

Kullanım:
    python main.py                                          # örnek sözleşmeyle 3 numaralı hakedişi hazırlar
    python main.py --sozlesme sozlesme.json --pozlar pozlar.xlsx --metraj metraj.xlsx --hakedis 4
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from collections import defaultdict
from dataclasses import dataclass
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from fractions import Fraction
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
KURUS = Decimal("0.01")
TEVKIFAT_ESIGI = Decimal("5000000")

VARSAYILAN = {
    "is_adi": "",
    "yuklenici": "",
    "isveren": "",
    "sozlesme_bedeli": None,          # verilmezse poz listesinden (sözleşme miktarı × birim fiyat)
    "kdv_orani": "0.20",
    "kdv_tevkifati": "otomatik",      # "otomatik" | "4/10" gibi bir pay | "yok"
    "belirlenmis_alici": False,       # kamu idaresi, belediye, KİT vb. (tebliğde sayılan alıcılar)
    "yillara_yaygin": True,           # iş takvim yılını aşıyorsa stopaj uygulanır
    "stopaj_orani": "0.05",
    "damga_vergisi_orani": "0",       # sözleşmenizde hakedişten kesileceği yazıyorsa (ör. 0.00948)
    "teminat_kesintisi_orani": "0",   # nakit teminat / tutulan kesinti oranı (ör. 0.05)
    "avans_tutari": "0",              # verilen avans; hakediş tutarı oranında mahsup edilir
    "fiyat_farki": {},                # {"hakediş no": tutar} — bu hakedişin fiyat farkı (hesaplanmış olarak)
    "diger_kesintiler": {},           # {"hakediş no": {"açıklama": tutar}}
}


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def para(x) -> Decimal:
    if x in (None, ""):
        return Decimal(0)
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("₺", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return Decimal(0)


def yz(x) -> str:
    """Sayıyı gereksiz sıfırlar olmadan Türkçe ondalıkla yazar: 20.00 → 20, 0.948 → 0,948."""
    return f"{float(x):g}".replace(".", ",")


def yuvarla(x: Decimal) -> Decimal:
    return x.quantize(KURUS, rounding=ROUND_HALF_UP)


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ayirici = csv.Sniffer().sniff(metin.splitlines()[0], delimiters=",;\t").delimiter
        satirlar = list(csv.reader(metin.splitlines(), delimiter=ayirici))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def sutun(baslik: list, *adlar: str) -> int | None:
    b = [kucuk(x) for x in baslik]
    return next((i for i, x in enumerate(b) if x in adlar), None)


@dataclass
class Poz:
    no: str
    tanim: str
    birim: str
    sozlesme_miktari: Decimal
    birim_fiyat: Decimal


def pozlari_oku(yol: Path) -> dict[str, Poz]:
    s = tablo_oku(yol)
    i_no = sutun(s[0], "poz no", "poz", "sıra no", "kalem no")
    i_t = sutun(s[0], "tanım", "iş kalemi", "imalat", "açıklama", "işin tanımı")
    i_b = sutun(s[0], "birim", "ölçü birimi")
    i_m = sutun(s[0], "sözleşme miktarı", "miktar", "teklif miktarı")
    i_f = sutun(s[0], "birim fiyat", "birim fiyatı", "teklif birim fiyatı")
    if None in (i_no, i_m, i_f):
        raise SystemExit(f"Poz listesinde Poz No, Sözleşme Miktarı ve Birim Fiyat sütunları gerekli. Başlıklar: {s[0]}")
    pozlar = {}
    for r in s[1:]:
        no = str(r[i_no] or "").strip()
        if no:
            pozlar[no] = Poz(no, str(r[i_t] or "").strip() if i_t is not None else "", str(r[i_b] or "").strip() if i_b is not None else "",
                             para(r[i_m]), para(r[i_f]))
    return pozlar


def metraj_oku(yol: Path, pozlar: dict[str, Poz]) -> tuple[dict[int, dict[str, Decimal]], list[str]]:
    """Hakediş no → poz → o dönemde yapılan miktar."""
    s = tablo_oku(yol)
    i_h = sutun(s[0], "hakediş no", "hakediş", "dönem", "hakedis no")
    i_p = sutun(s[0], "poz no", "poz")
    i_m = sutun(s[0], "miktar", "bu dönem miktar", "gerçekleşen miktar", "metraj")
    if None in (i_h, i_p, i_m):
        raise SystemExit(f"Metrajda Hakediş No, Poz No ve Miktar sütunları gerekli. Başlıklar: {s[0]}")
    metraj: dict[int, dict[str, Decimal]] = defaultdict(lambda: defaultdict(Decimal))
    uyarilar = []
    for r in s[1:]:
        p = str(r[i_p] or "").strip()
        try:
            h = int(para(r[i_h]))
        except (ValueError, InvalidOperation):
            continue
        if p not in pozlar:
            uyarilar.append(f"Metrajdaki {p} pozu sözleşme poz listesinde yok (yeni birim fiyat / onaylı ek poz gerekir); hesaba katılmadı")
            continue
        metraj[h][p] += para(r[i_m])
    return metraj, uyarilar


def ayarlari_oku(yol: Path | None) -> dict:
    a = dict(VARSAYILAN)
    if yol:
        a.update(json.loads(yol.read_text(encoding="utf-8")))
    return a


def oran(x) -> Decimal:
    if isinstance(x, str) and "/" in x:
        f = Fraction(x.strip())
        return Decimal(f.numerator) / Decimal(f.denominator)
    return para(x)


# ----------------------------------------------------------------------------
# Hesap
# ----------------------------------------------------------------------------

def icmal(pozlar: dict[str, Poz], metraj: dict[int, dict[str, Decimal]], a: dict, n: int) -> dict:
    """n numaralı hakediş için icmal. Tutarlar kümülatif hesaplanıp önceki hakedişten farkla bulunur."""
    sozlesme_bedeli = para(a["sozlesme_bedeli"]) if a.get("sozlesme_bedeli") not in (None, "") else \
        sum((p.sozlesme_miktari * p.birim_fiyat for p in pozlar.values()), Decimal(0))
    sozlesme_bedeli = yuvarla(sozlesme_bedeli)
    kdv_orani = oran(a["kdv_orani"])

    def kumulatif_is(k: int) -> Decimal:
        return yuvarla(sum((sum((metraj.get(h, {}).get(p, Decimal(0)) for h in range(1, k + 1)), Decimal(0)) * poz.birim_fiyat
                            for p, poz in pozlar.items()), Decimal(0)))

    def ff(k: int) -> Decimal:
        return yuvarla(sum((para(a["fiyat_farki"].get(str(h), 0)) for h in range(1, k + 1)), Decimal(0)))

    toplam_is, onceki_is = kumulatif_is(n), kumulatif_is(n - 1)
    toplam_ff, onceki_ff = ff(n), ff(n - 1)
    bu_is, bu_ff = toplam_is - onceki_is, toplam_ff - onceki_ff
    bu_hakedis = bu_is + bu_ff

    # KDV tevkifatı
    t = kucuk(a["kdv_tevkifati"])
    if t == "yok":
        tevkifat_orani, tevkifat_gerekce = Decimal(0), "Tevkifat uygulanmıyor (sözleşme ayarı)"
    elif t == "otomatik":
        kdv_dahil = sozlesme_bedeli * (1 + kdv_orani)
        if a.get("belirlenmis_alici"):
            tevkifat_orani, tevkifat_gerekce = Decimal("0.4"), "Belirlenmiş alıcı: 4/10 KDV tevkifatı"
        elif kdv_dahil >= TEVKIFAT_ESIGI:
            tevkifat_orani, tevkifat_gerekce = Decimal("0.4"), f"KDV dahil sözleşme bedeli {tl(kdv_dahil)} TL ≥ 5.000.000 TL: 4/10 KDV tevkifatı"
        else:
            tevkifat_orani, tevkifat_gerekce = Decimal(0), f"KDV dahil sözleşme bedeli {tl(kdv_dahil)} TL < 5.000.000 TL ve belirlenmiş alıcı değil: tevkifat yok"
    else:
        tevkifat_orani, tevkifat_gerekce = oran(a["kdv_tevkifati"]), f"Sözleşme ayarı: {a['kdv_tevkifati']} KDV tevkifatı"

    kdv = yuvarla(bu_hakedis * kdv_orani)
    tevkifat = yuvarla(kdv * tevkifat_orani)

    # Kesintiler
    stopaj_orani = oran(a["stopaj_orani"]) if a.get("yillara_yaygin") else Decimal(0)
    kesintiler = []
    if stopaj_orani:
        kesintiler.append((f"Gelir / kurumlar vergisi stopajı (%{yz(stopaj_orani * 100)})", yuvarla(bu_hakedis * stopaj_orani)))
    if oran(a["damga_vergisi_orani"]):
        d = oran(a["damga_vergisi_orani"])
        kesintiler.append((f"Damga vergisi (binde {yz(d * 1000)})", yuvarla(bu_hakedis * d)))
    if oran(a["teminat_kesintisi_orani"]):
        tk = oran(a["teminat_kesintisi_orani"])
        kesintiler.append((f"Teminat kesintisi (%{yz(tk * 100)})", yuvarla(bu_hakedis * tk)))
    avans = para(a["avans_tutari"])
    avans_mahsup = Decimal(0)
    if avans and sozlesme_bedeli:
        mahsup_orani = avans / sozlesme_bedeli
        onceki_mahsup = min(avans, yuvarla((onceki_is + onceki_ff) * mahsup_orani))
        avans_mahsup = min(avans - onceki_mahsup, yuvarla(bu_hakedis * mahsup_orani))
        kesintiler.append((f"Avans mahsubu (avans / sözleşme bedeli = %{mahsup_orani * 100:.2f})", avans_mahsup))
    for ad, tutar in (a["diger_kesintiler"].get(str(n)) or {}).items():
        kesintiler.append((ad, yuvarla(para(tutar))))
    kesinti_toplami = sum((k[1] for k in kesintiler), Decimal(0))

    tahakkuk = bu_hakedis + kdv - tevkifat
    odenecek = tahakkuk - kesinti_toplami
    return {
        "n": n, "sozlesme_bedeli": sozlesme_bedeli, "toplam_is": toplam_is, "onceki_is": onceki_is, "bu_is": bu_is,
        "toplam_ff": toplam_ff, "onceki_ff": onceki_ff, "bu_ff": bu_ff, "bu_hakedis": bu_hakedis,
        "kdv_orani": kdv_orani, "kdv": kdv, "tevkifat_orani": tevkifat_orani, "tevkifat": tevkifat, "tevkifat_gerekce": tevkifat_gerekce,
        "kesintiler": kesintiler, "kesinti_toplami": kesinti_toplami, "tahakkuk": tahakkuk, "odenecek": odenecek,
        "avans_mahsup": avans_mahsup, "ilerleme": None if not sozlesme_bedeli else toplam_is / sozlesme_bedeli,
    }


def poz_ozeti(pozlar, metraj, n):
    satirlar = []
    for p, poz in pozlar.items():
        onceki = sum((metraj.get(h, {}).get(p, Decimal(0)) for h in range(1, n)), Decimal(0))
        bu = metraj.get(n, {}).get(p, Decimal(0))
        toplam = onceki + bu
        asim = toplam - poz.sozlesme_miktari if toplam > poz.sozlesme_miktari else Decimal(0)
        satirlar.append({"poz": poz, "onceki": onceki, "bu": bu, "toplam": toplam, "tutar": yuvarla(toplam * poz.birim_fiyat),
                         "bu_tutar": yuvarla(bu * poz.birim_fiyat),
                         "oran": None if not poz.sozlesme_miktari else toplam / poz.sozlesme_miktari, "asim": asim})
    return satirlar


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
PARA = "#,##0.00"
MIKTAR = "#,##0.000"


def tl(x) -> str:
    return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def calistir(sozlesme: Path | None, pozlar_yolu: Path, metraj_yolu: Path, n: int | None, cikti: Path) -> dict:
    a = ayarlari_oku(sozlesme)
    pozlar = pozlari_oku(pozlar_yolu)
    metraj, uyarilar = metraj_oku(metraj_yolu, pozlar)
    if not metraj:
        raise SystemExit("Metraj dosyasında hakediş kaydı yok.")
    n = n or max(metraj)
    ic = icmal(pozlar, metraj, a, n)
    ozet = poz_ozeti(pozlar, metraj, n)
    for x in ozet:
        if x["asim"] > 0:
            uyarilar.append(f"Poz {x['poz'].no}: kümülatif miktar sözleşme miktarını {yz(x['asim'])} {x['poz'].birim} aşıyor "
                            "(iş artışı / onaylı fazla metraj gerekir)")
    if ic["ilerleme"] is not None and ic["ilerleme"] > 1:
        uyarilar.append(f"Kümülatif iş tutarı sözleşme bedelinin %{ic['ilerleme'] * 100:.1f}'i: iş artışı onayı ve sınırlar "
                        "(kamu işlerinde 4735 sayılı Kanun md. 24) kontrol edilmeli")

    wb = Workbook()
    ws = wb.active
    ws.title = "Hakediş İcmali"
    ws.append([f"{n} NO'LU HAKEDİŞ İCMALİ"])
    ws["A1"].font = Font(bold=True, size=13)
    for etiket, deger in (("İşin adı", a["is_adi"]), ("Yüklenici", a["yuklenici"]), ("İşveren", a["isveren"]),
                          ("Sözleşme bedeli (KDV hariç)", float(ic["sozlesme_bedeli"])),
                          ("Fiziki ilerleme (iş tutarı / sözleşme bedeli)", None if ic["ilerleme"] is None else float(ic["ilerleme"]))):
        ws.append([etiket, deger])
    ws["B5"].number_format = PARA
    ws["B6"].number_format = "0.00%"
    ws.append([])
    ws.append(["", "Toplam (kümülatif)", "Önceki hakedişler", "Bu hakediş"])
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    ws.append(["A. Sözleşme fiyatlarıyla yapılan iş", float(ic["toplam_is"]), float(ic["onceki_is"]), float(ic["bu_is"])])
    ws.append(["B. Fiyat farkı", float(ic["toplam_ff"]), float(ic["onceki_ff"]), float(ic["bu_ff"])])
    ws.append(["C. Hakediş tutarı (A + B)", float(ic["toplam_is"] + ic["toplam_ff"]), float(ic["onceki_is"] + ic["onceki_ff"]),
               float(ic["bu_hakedis"])])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for r in range(ws.max_row - 2, ws.max_row + 1):
        for c in (2, 3, 4):
            ws.cell(r, c).number_format = PARA
    ws.append([])
    satirlar = [(f"D. KDV (%{yz(ic['kdv_orani'] * 100)})", ic["kdv"]),
                (f"E. KDV tevkifatı ({_pay(ic['tevkifat_orani'])}) — alıcı tarafından beyan edilir", -ic["tevkifat"]),
                ("F. Tahakkuk (C + D − E)", ic["tahakkuk"])]
    satirlar += [(f"   − {ad}", -t) for ad, t in ic["kesintiler"]]
    satirlar += [("G. Kesintiler toplamı", -ic["kesinti_toplami"]), ("ÖDENECEK NET TUTAR (F − G)", ic["odenecek"])]
    ws.append(["Bu hakediş", "Tutar"])
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    ws.append(["C. Hakediş tutarı", float(ic["bu_hakedis"])])
    ws.cell(ws.max_row, 2).number_format = PARA
    for etiket, t in satirlar:
        ws.append([etiket, float(t)])
        ws.cell(ws.max_row, 2).number_format = PARA
        if etiket.startswith(("F.", "ÖDENECEK")):
            ws.cell(ws.max_row, 1).font = ws.cell(ws.max_row, 2).font = Font(bold=True)
    ws.append([])
    ws.append(["Tevkifat gerekçesi", ic["tevkifat_gerekce"]])
    ws.append(["Fatura", f"Yüklenici faturası: matrah {tl(ic['bu_hakedis'])} TL + KDV {tl(ic['kdv'])} TL; "
                         f"tevkif edilen KDV {tl(ic['tevkifat'])} TL (KDV dahil {tl(ic['bu_hakedis'] + ic['kdv'])} TL)"])
    for u in uyarilar:
        ws.append(["Uyarı", u])
    ws.column_dimensions["A"].width = 62
    for c in "BCD":
        ws.column_dimensions[c].width = 20

    m = wb.create_sheet("Metraj Özeti")
    m.append(["Poz No", "Tanım", "Birim", "Sözleşme Miktarı", "Birim Fiyat", "Önceki Miktar", "Bu Dönem Miktar",
              "Toplam Miktar", "Gerçekleşme %", "Bu Dönem Tutar", "Toplam Tutar", "Sözleşme Aşımı"])
    for h in m[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for x in ozet:
        p = x["poz"]
        m.append([p.no, p.tanim, p.birim, float(p.sozlesme_miktari), float(p.birim_fiyat), float(x["onceki"]), float(x["bu"]),
                  float(x["toplam"]), None if x["oran"] is None else float(x["oran"]), float(x["bu_tutar"]), float(x["tutar"]),
                  float(x["asim"]) if x["asim"] else None])
        r = m.max_row
        for c in (4, 6, 7, 8, 12):
            m.cell(r, c).number_format = MIKTAR
        for c in (5, 10, 11):
            m.cell(r, c).number_format = PARA
        m.cell(r, 9).number_format = "0.0%"
        if x["asim"]:
            m.cell(r, 12).fill = KIRMIZI
    m.append(["", "TOPLAM", "", "", "", "", "", "", "", float(sum((x["bu_tutar"] for x in ozet), Decimal(0))),
              float(sum((x["tutar"] for x in ozet), Decimal(0)))])
    for c in (10, 11):
        m.cell(m.max_row, c).number_format = PARA
        m.cell(m.max_row, c).font = Font(bold=True)
    for j, w in enumerate((12, 42, 8, 15, 14, 14, 15, 14, 13, 16, 16, 14), 1):
        m.column_dimensions[get_column_letter(j)].width = w
    m.freeze_panes = "C2"

    g = wb.create_sheet("Hakediş Geçmişi")
    g.append(["Hakediş No", "Hakediş Tutarı (C)", "KDV", "KDV Tevkifatı", "Kesintiler", "Ödenecek", "Kümülatif İş Tutarı", "İlerleme"])
    for h in g[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for k in range(1, n + 1):
        x = icmal(pozlar, metraj, a, k)
        g.append([k, float(x["bu_hakedis"]), float(x["kdv"]), float(x["tevkifat"]), float(x["kesinti_toplami"]), float(x["odenecek"]),
                  float(x["toplam_is"]), None if x["ilerleme"] is None else float(x["ilerleme"])])
        for c in range(2, 8):
            g.cell(g.max_row, c).number_format = PARA
        g.cell(g.max_row, 8).number_format = "0.0%"
    for j in range(1, 9):
        g.column_dimensions[get_column_letter(j)].width = 18

    b = wb.create_sheet("Bilgi")
    for s in [["Yöntem", "Tutarlar kümülatif hesaplanır; bu hakediş = toplam − önceki hakedişler (yuvarlama farkı birikmez)"],
              ["KDV tevkifatı", "Yapım işlerinde 4/10: belirlenmiş alıcılar veya KDV dahil bedeli 5.000.000 TL ve üzeri işler "
                                "(KDV Genel Uygulama Tebliği I/C-2.1.3.2.1, Seri No: 35, 01.03.2021)"],
              ["Stopaj", "Yıllara yaygın inşaat ve onarım işlerinde hakediş tutarı üzerinden %5 (GVK 94/3, KVK 15; 3491 sayılı CBK). "
                         "Demiryolu, gemi ve nükleer santral inşaatında %1 — sozlesme.json'da stopaj_orani ile değiştirin"],
              ["Avans mahsubu", "Bu hakediş tutarı × (avans / sözleşme bedeli); toplam mahsup avansı aşmaz"],
              ["Fiyat farkı", "Bu paket fiyat farkını hesaplamaz; hesaplanmış tutar sozlesme.json'daki fiyat_farki alanından alınır"],
              ["Uyarı", "Sonuçlar sözleşme şartlarına ve güncel mevzuata göre mali müşavir / kontrol mühendisi tarafından teyit edilmelidir"]]:
        b.append(s)
    b.column_dimensions["A"].width = 16
    b.column_dimensions["B"].width = 130
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    return {"icmal": ic, "ozet": ozet, "uyarilar": uyarilar}


def _pay(x: Decimal) -> str:
    if not x:
        return "yok"
    f = Fraction(str(x)).limit_denominator(10)
    return f"{f.numerator * 10 // f.denominator}/10" if 10 % f.denominator == 0 else str(f)


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description="Poz listesi ve dönem metrajlarından hakediş icmali hazırlar.")
    ap.add_argument("--sozlesme", type=Path, default=BURASI / "ornek_veri" / "sozlesme.json", help="Sözleşme ayarları (JSON)")
    ap.add_argument("--pozlar", type=Path, default=BURASI / "ornek_veri" / "pozlar.csv",
                    help="Poz listesi (.xlsx/.csv): Poz No, Tanım, Birim, Sözleşme Miktarı, Birim Fiyat")
    ap.add_argument("--metraj", type=Path, default=BURASI / "ornek_veri" / "metraj.csv",
                    help="Dönem metrajları (.xlsx/.csv): Hakediş No, Poz No, Miktar")
    ap.add_argument("--hakedis", type=int, help="Hazırlanacak hakediş no (varsayılan: metrajdaki son hakediş)")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "hakedis.xlsx")
    a = ap.parse_args(argv)
    s = calistir(a.sozlesme, a.pozlar, a.metraj, a.hakedis, a.cikti)
    ic = s["icmal"]
    print(f"[OK] {ic['n']} no'lu hakediş · bu hakediş {tl(ic['bu_hakedis'])} TL · KDV {tl(ic['kdv'])} · "
          f"tevkifat {tl(ic['tevkifat'])} · kesintiler {tl(ic['kesinti_toplami'])}")
    print(f"[OK] Ödenecek net tutar: {tl(ic['odenecek'])} TL · ilerleme %{yz(round((ic['ilerleme'] or 0) * 100, 1))}")
    for u in s["uyarilar"]:
        print(f"[!] {u}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
