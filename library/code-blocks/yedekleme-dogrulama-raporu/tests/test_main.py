import sys
import tempfile
import unittest
from datetime import datetime, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
AN = datetime(2026, 10, 12, 9, 0)


def k(ad, saat_once, durum="Başarılı", boyut=None, sure=60, tur="Tam"):
    bas = AN - timedelta(hours=saat_once)
    return {"is": ad, "sistem": "S", "bas": bas, "bit": bas + timedelta(minutes=sure) if sure is not None else None, "durum": durum, "boyut": None if boyut is None else D(boyut), "tur": tur}


def p(ad, aralik=24, azami=None):
    return {"is": ad, "sistem": "S", "aralik": None if aralik is None else D(aralik), "azami": None if azami is None else D(azami)}


def bul(r, ad):
    return next(i for i in r["isler"] if i["is"] == ad)


class DogrulamaTesti(unittest.TestCase):
    def test_gecikme_ve_durum(self):
        ks = [k("sınırda", 31), k("gecikmiş", 31.5), k("başarısız-son", 50), k("başarısız-son", 2, "Başarısız"), k("hiç-geçerli-yok", 5, "Başarısız"), k("toparlanmış", 30, "Başarısız"), k("toparlanmış", 6),
              k("uyarılı", 4, "Uyarı"), k("plansız", 500), k("bitişsiz", 10, sure=None), k("gelecek", -5)]
        r = main.dogrula(ks, [p(x) for x in ("sınırda", "gecikmiş", "başarısız-son", "hiç-geçerli-yok", "toparlanmış", "uyarılı", "hiç-çalışmamış", "bitişsiz", "gelecek")], AN)
        self.assertEqual({i["is"]: (i["durum"], i.get("yas"), i["toplam"], i["basarili"], i["basarisiz"]) for i in r["isler"]},
                         {"sınırda": ("Tamam", D("30.0"), 1, 1, 0), "gecikmiş": ("Sorunlu", D("30.5"), 1, 1, 0), "başarısız-son": ("Sorunlu", D("49.0"), 2, 1, 1), "hiç-geçerli-yok": ("Sorunlu", None, 1, 0, 1),
                          "toparlanmış": ("Dikkat", D("5.0"), 2, 1, 1), "uyarılı": ("Dikkat", D("3.0"), 1, 0, 0), "plansız": ("Tamam", D("499.0"), 0, 0, 0), "bitişsiz": ("Tamam", D("10.0"), 1, 1, 0),
                          "hiç-çalışmamış": ("Sorunlu", None, 0, 0, 0), "gelecek": ("Sorunlu", None, 0, 0, 0)})
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Bilgi", "Planda olmayan iş", "plansız"), ("Orta", "Pencerede başarısız koşu", "toparlanmış"), ("Orta", "Uyarıyla tamamlandı", "uyarılı"), ("Yüksek", "Geçerli yedek yok", "hiç-geçerli-yok"),
                          ("Yüksek", "Planlı iş hiç çalışmamış", "gelecek"), ("Yüksek", "Planlı iş hiç çalışmamış", "hiç-çalışmamış"), ("Yüksek", "Son koşu başarısız", "başarısız-son"), ("Yüksek", "Son koşu başarısız", "hiç-geçerli-yok"),
                          ("Yüksek", "Yedek gecikmiş", "başarısız-son"), ("Yüksek", "Yedek gecikmiş", "gecikmiş")])
        self.assertEqual((r["say"], [i["is"] for i in r["isler"]][:5], [(x["is"], x["durum"]) for x in r["sorunlu"]], bul(r, "uyarılı")["oran"], bul(r, "toparlanmış")["oran"], bul(r, "plansız")["oran"]),
                         ({"Sorunlu": 5, "Dikkat": 2, "Tamam": 3}, ["başarısız-son", "gecikmiş", "gelecek", "hiç-çalışmamış", "hiç-geçerli-yok"],
                          [("başarısız-son", "Başarısız"), ("uyarılı", "Uyarı"), ("hiç-geçerli-yok", "Başarısız"), ("toparlanmış", "Başarısız")], D("0.00"), D("50.00"), None))
        self.assertEqual(bul(main.dogrula([k("gecikmiş", 31.5)], [p("gecikmiş")], AN, 7, D(8)), "gecikmiş")["durum"], "Tamam")
        self.assertEqual(bul(main.dogrula([k("x", 200)], [p("x", 168)], AN, 7), "x")["toplam"], 0)

    def test_boyut_ve_sure(self):
        def son(boyutlar, tur_son="Tam", **kw):
            ks = [k("iş", 24 * (len(boyutlar) - n), boyut=b, tur="Tam") for n, b in enumerate(boyutlar[:-1])] + [k("iş", 2, boyut=boyutlar[-1], tur=tur_son, **kw)]
            r = main.dogrula(ks, [p("iş", 24, 90)], AN, 30)
            return bul(r, "iş")["durum"], [u["tur"] for u in r["uyarilar"]]
        self.assertEqual(son([100, 110, 120, 55]), ("Dikkat", ["Yedek boyutu olağandışı küçük"]))                 # medyan 110 → sınır 55
        self.assertEqual(son([100, 110, 120, 56]), ("Tamam", []))
        self.assertEqual(son([100, 110, 120, 130, 57.5]), ("Dikkat", ["Yedek boyutu olağandışı küçük"]))          # medyan 115 → sınır 57,5
        self.assertEqual(son([100, 110, 10]), ("Tamam", []))                                                      # en az 3 önceki yedek gerekir
        self.assertEqual(son([100, 110, 120, 10], "Artımlı"), ("Tamam", []))                                      # farklı türle karşılaştırılmaz
        self.assertEqual(son([100, 110, 120, 0]), ("Sorunlu", ["Yedek boyutu sıfır"]))
        self.assertEqual(son([100, 110, 120, 115], sure=91), ("Dikkat", ["Süre aşımı"]))
        self.assertEqual(son([100, 110, 120, 115], sure=90), ("Tamam", []))
        self.assertEqual((main.medyan([D(3), D(1), D(2)]), main.medyan([D(4), D(1), D(2), D(3)])), (D(2), D("2.5")))
        with self.assertRaises(ValueError):
            main.dogrula([], [], AN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "y.xlsx"
            r = main.calistir(ORNEK / "gecmis.csv", ORNEK / "plan.csv", cikti, AN)
            self.assertEqual([(i["is"], i["durum"], i.get("yas"), i["toplam"], i["basarili"], i["basarisiz"], i["oran"]) for i in r["isler"]],
                             [("Muhasebe Veritabanı", "Sorunlu", D("55.7"), 7, 5, 2, D("71.43")), ("Web Sunucusu", "Sorunlu", None, 0, 0, 0, None), ("E-posta Sunucusu", "Dikkat", D("5.6"), 7, 7, 0, D("100.00")),
                              ("Sanal Makineler Haftalık", "Dikkat", D("1.5"), 1, 0, 0, D("0.00")), ("Arşiv Aylık", "Tamam", D("267.0"), 0, 0, 0, None), ("Dosya Sunucusu Günlük", "Tamam", D("6.0"), 7, 7, 0, D("100.00")),
                              ("Test Ortamı", "Tamam", D("68.8"), 1, 1, 0, D("100.00"))])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Yüksek", "Planlı iş hiç çalışmamış", "Web Sunucusu"), ("Yüksek", "Son koşu başarısız", "Muhasebe Veritabanı"), ("Yüksek", "Yedek gecikmiş", "Muhasebe Veritabanı"),
                              ("Orta", "Süre aşımı", "Sanal Makineler Haftalık"), ("Orta", "Uyarıyla tamamlandı", "Sanal Makineler Haftalık"), ("Orta", "Yedek boyutu olağandışı küçük", "E-posta Sunucusu"),
                              ("Bilgi", "Planda olmayan iş", "Test Ortamı")])
            self.assertEqual((len(r["sorunlu"]), bul(r, "Sanal Makineler Haftalık")["son"]["sure"], "121.5 GB" in next(u["aciklama"] for u in r["uyarilar"] if u["tur"].startswith("Yedek boyutu"))), (3, 570, True))
            self.assertEqual(load_workbook(cikti).sheetnames, ["İş Özeti", "Sorunlu Koşular", "Uyarılar"])

    def test_okuma_ve_cli(self):
        self.assertEqual([main.zaman(x) for x in ("12.10.2026 09:05", "2026-10-12T09:05:30", "12.10.2026", datetime(2026, 1, 2, 3, 4), "x", None)],
                         [datetime(2026, 10, 12, 9, 5), datetime(2026, 10, 12, 9, 5, 30), datetime(2026, 10, 12), datetime(2026, 1, 2, 3, 4), None, None])
        self.assertEqual([main.olcu(x) for x in ("12.5", "12,5", "1.250,5", 7, "", "x")], [D("12.5"), D("12.5"), D("1250.5"), D(7), None, None])
        with tempfile.TemporaryDirectory() as tmp:
            y, pl = Path(tmp) / "g.csv", Path(tmp) / "p.csv"
            y.write_text("Job Name;Server;Start Time;End Time;Status;Size (GB)\nA;S1;2026-10-12 01:00:00;2026-10-12 01:30:00;Success;12.5\nA;S1;2026-10-11 01:00;2026-10-11 00:30;WARNING;\nB;S2;x;;Failed;\nB;S2;11.10.2026 02:00;;Skipped;\n",
                         encoding="utf-8")
            pl.write_text("İş;Aralık (saat);Süre Sınırı (dk)\nA;12;45\nB;Daily;\nC;bazen;\n", encoding="utf-8")
            l, uy = main.log_oku(y)
            self.assertEqual(([(x["is"], x["bas"], x["bit"], x["durum"], x["boyut"]) for x in l], [u["tur"] for u in uy]),
                             ([("A", datetime(2026, 10, 12, 1), datetime(2026, 10, 12, 1, 30), "Başarılı", D("12.5")), ("A", datetime(2026, 10, 11, 1), None, "Uyarı", None), ("B", datetime(2026, 10, 11, 2), None, "Bilinmiyor", None)],
                              ["Koşu okunamadı", "Durum tanınmadı"]))
            pln, uy = main.plan_oku(pl)
            self.assertEqual(([(x["is"], x["aralik"], x["azami"]) for x in pln], [u["tur"] for u in uy]), ([("A", D(12), D(45)), ("B", D(24), None), ("C", None, None)], ["Sıklık okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--gecmis", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--plan", str(pl), "--an", "12.10.2026 09:00", "--gun", "14", "--tolerans", "12", "--boyut-sapma", "30"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + x) for x in (["--an", "x"], ["--gun", "0"], ["--tolerans", "-1"], ["--boyut-sapma", "100"])], [2] * 4)
            self.assertEqual((main.main(["--gecmis", str(Path(tmp) / "yok.csv")]), main.main(arg + ["--plan", str(Path(tmp) / "yok.csv")])), (1, 1))
            for icerik in ("İş Adı;Durum\nA;x\n", "İş Adı;Başlangıç;Durum\n;;\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
