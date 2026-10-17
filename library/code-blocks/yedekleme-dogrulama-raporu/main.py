"""
Yedekleme Doğrulama Raporu — Gerisi Bende kod bloğu
Bilgi Teknolojileri › Sistem Uzmanı

Yedekleme yazılımının iş geçmişinden başarısız, eksik, geciken ve boyutu şüpheli yedekleri raporlar.
  Geçerli yedek       = durumu "Başarılı" veya "Uyarı" olan koşu (uyarılı koşular ayrıca işaretlenir)
  Son yedek yaşı      = --an − son geçerli yedeğin bitiş saati (saat)
  Gecikmiş            = yaş > beklenen aralık + --tolerans (6 saat);  aralık: Günlük 24, Haftalık 168, Aylık 744 saat veya plandaki sayı
  Pencere             = --an anından geriye --gun (7) gün; başarı oranı = başarılı koşu ÷ penceredeki koşu
  Boyut şüphesi       = son geçerli yedeğin boyutu, önceki geçerli yedeklerin (en az 3) medyanından %--boyut-sapma (50) veya daha fazla küçük
  Süre aşımı          = son koşunun süresi plandaki "Azami Süre (dk)" değerinden uzun
  Eksik               = planda olup geçmişte hiç koşusu olmayan iş
"Başarılı" yazan bir yedek geri yüklenebilir demek değildir; düzenli geri yükleme testi yapın.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
LOG = {"is": ("is adi", "is", "gorev", "job", "job name", "yedekleme isi"), "sistem": ("sistem", "kaynak", "sunucu", "server"), "bas": ("baslangic", "baslangic zamani", "start time", "start"),
       "bit": ("bitis", "bitis zamani", "end time", "end"), "durum": ("durum", "sonuc", "status", "result"), "boyut": ("boyut", "boyut gb", "size", "size gb"), "tur": ("tur", "yedek turu", "type")}
PLAN = {"is": ("is adi", "is", "gorev", "job", "job name", "yedekleme isi"), "sistem": ("sistem", "kaynak", "sunucu", "server"), "aralik": ("siklik", "beklenen aralik", "beklenen aralik saat", "aralik", "aralik saat"),
        "azami": ("azami sure", "azami sure dk", "sure siniri", "sure siniri dk")}
DURUM = {"basarili": "Başarılı", "success": "Başarılı", "successful": "Başarılı", "completed": "Başarılı", "ok": "Başarılı", "tamamlandi": "Başarılı", "uyari": "Uyarı", "warning": "Uyarı", "uyarili": "Uyarı",
         "completed with warnings": "Uyarı", "basarisiz": "Başarısız", "failed": "Başarısız", "fail": "Başarısız", "error": "Başarısız", "hata": "Başarısız", "iptal": "Başarısız", "cancelled": "Başarısız"}
SIKLIK = {"saatlik": 1, "gunluk": 24, "her gun": 24, "daily": 24, "haftalik": 168, "weekly": 168, "aylik": 744, "monthly": 744}
GECERLI = ("Başarılı", "Uyarı")
S_RENK = {"Sorunlu": KIRMIZI, "Dikkat": SARI, "Tamam": YESIL}
ZF = "DD.MM.YYYY HH:MM"


def zaman(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    if isinstance(x, date):
        return datetime(x.year, x.month, x.day)
    s = str(x or "").strip().replace("T", " ")[:19]
    for f in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d/%m/%Y %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def olcu(x) -> Decimal | None:
    """Boyut ve süre gibi ölçüler: virgül veya tek nokta ondalık ayırıcıdır (12.5 = 12,5)."""
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x or ""))
    s = s.replace(".", "").replace(",", ".") if "," in s else s
    try:
        return Decimal(s) if s else None
    except InvalidOperation:
        return None


def medyan(l: list[Decimal]) -> Decimal:
    s, n = sorted(l), len(l)
    return s[n // 2] if n % 2 else (s[n // 2 - 1] + s[n // 2]) / 2


def log_oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, LOG, ("is", "bas", "durum")):
        ad = metin(r["is"])
        if not ad:
            continue
        bas, d = zaman(r["bas"]), DURUM.get(katla(r["durum"]))
        if bas is None:
            uy.append(uyari("Yüksek", "Koşu okunamadı", f"{ad} (satır {r['_satir']})", f"Başlangıç '{metin(r['bas'])}' okunamadı (GG.AA.YYYY SS:DD). Koşu rapora alınmadı."))
            continue
        if d is None:
            uy.append(uyari("Orta", "Durum tanınmadı", f"{ad} · {bas:%d.%m.%Y %H:%M}", f"'{metin(r['durum'])}' durumu tanınmadı; koşu geçerli yedek sayılmadı."))
        bit = zaman(r["bit"])
        l.append({"is": ad, "sistem": metin(r["sistem"]), "bas": bas, "bit": bit if bit and bit >= bas else None, "durum": d or "Bilinmiyor", "boyut": olcu(r["boyut"]), "tur": metin(r["tur"])})
    return l, uy


def plan_oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, PLAN, ("is", "aralik")):
        ad = metin(r["is"])
        if not ad:
            continue
        a = SIKLIK.get(katla(r["aralik"])) or olcu(r["aralik"])
        if not a or a <= 0:
            uy.append(uyari("Orta", "Sıklık okunamadı", ad, f"'{metin(r['aralik'])}' okunamadı; Günlük, Haftalık, Aylık veya saat cinsinden sayı yazın. İş için gecikme denetimi yapılmadı."))
            a = None
        az = olcu(r["azami"])
        l.append({"is": ad, "sistem": metin(r["sistem"]), "aralik": Decimal(a) if a else None, "azami": az if az and az > 0 else None})
    return l, uy


def dogrula(kosular: list[dict], plan: list[dict], an: datetime, gun: int = 7, tolerans: Decimal = Decimal(6), sapma: Decimal = Decimal(50)) -> dict:
    if not kosular and not plan:
        raise ValueError("Yedekleme geçmişi ve plan boş.")
    uy, isler, pencere_bas = [], {}, an - timedelta(days=gun)
    for p in plan:
        isler[katla(p["is"])] = {"is": p["is"], "sistem": p["sistem"], "aralik": p["aralik"], "azami": p["azami"], "planda": True, "kosular": []}
    for k in sorted(kosular, key=lambda k: k["bas"]):
        if k["bas"] > an:
            continue
        k["son_an"] = k["bit"] or k["bas"]
        k["sure"] = int((k["bit"] - k["bas"]).total_seconds() // 60) if k["bit"] else None
        i = isler.setdefault(katla(k["is"]), {"is": k["is"], "sistem": k["sistem"], "aralik": None, "azami": None, "planda": False, "kosular": []})
        i["sistem"] = i["sistem"] or k["sistem"]
        i["kosular"].append(k)
    sorunlu = []
    for i in isler.values():
        ks, onem = i["kosular"], set()

        def ekle(o, tur, aciklama, i=i, onem=onem):
            onem.add(o)
            uy.append(uyari(o, tur, i["is"], aciklama))

        pen = [k for k in ks if k["son_an"] >= pencere_bas]
        gecerli = [k for k in ks if k["durum"] in GECERLI]
        i["toplam"], i["basarili"], i["uyarili"], i["basarisiz"] = len(pen), sum(k["durum"] == "Başarılı" for k in pen), sum(k["durum"] == "Uyarı" for k in pen), sum(k["durum"] not in GECERLI for k in pen)
        i["oran"] = yuvarla(Decimal(100 * i["basarili"]) / len(pen)) if pen else None
        i["son"] = ks[-1] if ks else None
        i["son_gecerli"] = gecerli[-1]["son_an"] if gecerli else None
        i["yas"] = yuvarla(Decimal(int((an - i["son_gecerli"]).total_seconds())) / 3600, "0.1") if gecerli else None
        i["son_boyut"] = gecerli[-1]["boyut"] if gecerli else None
        sorunlu += [k for k in pen if k["durum"] != "Başarılı"]
        if not i["planda"]:
            ekle("Bilgi", "Planda olmayan iş", "Geçmişte koşusu var ama planda yok; beklenen aralık bilinmediği için gecikme denetimi yapılmadı.")
        if not ks:
            ekle("Yüksek", "Planlı iş hiç çalışmamış", f"{i['sistem'] or 'Sistem'} için planlı yedekleme işinin geçmişte hiçbir koşusu yok.")
            i["durum"] = "Sorunlu"
            continue
        if not gecerli:
            ekle("Yüksek", "Geçerli yedek yok", f"Geçmişteki {len(ks)} koşunun hiçbiri başarılı değil.")
        elif i["aralik"] and i["yas"] > i["aralik"] + tolerans:
            ekle("Yüksek", "Yedek gecikmiş", f"Son geçerli yedek {i['son_gecerli']:%d.%m.%Y %H:%M}, {i['yas']} saat önce; beklenen aralık {i['aralik']} saat (+{tolerans} saat tolerans).")
        if i["son"]["durum"] not in GECERLI:
            ekle("Yüksek", "Son koşu başarısız", f"{i['son']['bas']:%d.%m.%Y %H:%M} koşusu '{i['son']['durum']}'. Pencerede {i['basarisiz']} başarısız koşu var.")
        elif i["basarisiz"]:
            ekle("Orta", "Pencerede başarısız koşu", f"Son {gun} günde {i['basarisiz']} başarısız koşu var; son koşu geçerli.")
        if i["son"]["durum"] == "Uyarı":
            ekle("Orta", "Uyarıyla tamamlandı", f"{i['son']['bas']:%d.%m.%Y %H:%M} koşusu uyarıyla bitti; atlanan dosya veya disk olup olmadığına yedekleme yazılımından bakın.")
        if gecerli:
            turdes = [k["boyut"] for k in gecerli[:-1] if k["boyut"] is not None and katla(k["tur"]) == katla(gecerli[-1]["tur"])]
            if i["son_boyut"] is not None and i["son_boyut"] <= 0:
                ekle("Yüksek", "Yedek boyutu sıfır", f"{gecerli[-1]['bas']:%d.%m.%Y %H:%M} koşusu geçerli görünüyor ama boyutu sıfır.")
            elif i["son_boyut"] is not None and len(turdes) >= 3 and i["son_boyut"] <= medyan(turdes) * (100 - sapma) / 100:
                ekle("Orta", "Yedek boyutu olağandışı küçük", f"Son boyut {i['son_boyut']} GB; önceki {len(turdes)} geçerli yedeğin medyanı {medyan(turdes).normalize():f} GB. Kaynağın eksik yedeklenmediğini doğrulayın.")
        if i["azami"] and i["son"]["sure"] is not None and i["son"]["sure"] > i["azami"]:
            ekle("Orta", "Süre aşımı", f"Son koşu {i['son']['sure']} dakika sürdü; plandaki sınır {i['azami']} dakika.")
        i["durum"] = "Sorunlu" if "Yüksek" in onem else "Dikkat" if "Orta" in onem else "Tamam"
    sira = {"Sorunlu": 0, "Dikkat": 1, "Tamam": 2}
    liste = sorted(isler.values(), key=lambda i: (sira[i["durum"]], katla(i["is"])))
    return {"isler": liste, "sorunlu": sorted(sorunlu, key=lambda k: (k["bas"], katla(k["is"])), reverse=True), "an": an, "pencere_bas": pencere_bas, "say": {d: sum(i["durum"] == d for i in liste) for d in sira}, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "İş Özeti"
    _baslik(ws, ["Durum", "İş Adı", "Sistem", "Beklenen Aralık (saat)", "Son Geçerli Yedek", "Yaş (saat)", "Son Koşu Durumu", "Son Boyut (GB)", "Son Süre (dk)", "Pencerede Koşu", "Başarılı", "Uyarılı", "Başarısız", "Başarı %"],
            (10, 30, 14, 12, 17, 10, 14, 12, 10, 10, 9, 9, 10, 10))
    for i in r["isler"]:
        n = _satir(ws, [i["durum"], i["is"], i["sistem"], i["aralik"], i.get("son_gecerli"), i.get("yas"), i["son"]["durum"] if i.get("son") else "Koşu yok", i.get("son_boyut"), i["son"]["sure"] if i.get("son") else None, i["toplam"],
                        i["basarili"], i["uyarili"] or None, i["basarisiz"] or None, i["oran"]], para_s=(14,), dolgu={1: S_RENK[i["durum"]], 13: KIRMIZI if i["basarisiz"] else None})
        ws.cell(n, 5).number_format = ZF
    ws.append([])
    ws.append([f"Rapor anı {r['an']:%d.%m.%Y %H:%M} · pencere başlangıcı {r['pencere_bas']:%d.%m.%Y %H:%M}. 'Başarılı' yazan yedek geri yüklenebilir demek değildir; geri yükleme testi ayrıca yapılmalıdır."])
    ws = wb.create_sheet("Sorunlu Koşular")
    _baslik(ws, ["İş Adı", "Sistem", "Başlangıç", "Bitiş", "Durum", "Tür", "Boyut (GB)", "Süre (dk)"], (30, 14, 17, 17, 12, 12, 12, 10))
    for k in r["sorunlu"]:
        n = _satir(ws, [k["is"], k["sistem"], k["bas"], k["bit"], k["durum"], k["tur"], k["boyut"], k["sure"]], dolgu={5: SARI if k["durum"] == "Uyarı" else KIRMIZI})
        ws.cell(n, 3).number_format = ws.cell(n, 4).number_format = ZF
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(log_yol: Path, plan_yol: Path | None, cikti: Path, an: datetime, gun: int = 7, tolerans: Decimal = Decimal(6), sapma: Decimal = Decimal(50)) -> dict:
    kosular, uy = log_oku(log_yol)
    plan = []
    if plan_yol:
        plan, u = plan_oku(plan_yol)
        uy += u
    r = dogrula(kosular, plan, an, gun, tolerans, sapma)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Yedekleme iş geçmişinden başarısız, eksik ve geciken yedekleri raporlar.")
    p.add_argument("--gecmis", type=Path, help="İş Adı, Sistem, Başlangıç, Bitiş, Durum, Boyut (GB), Tür")
    p.add_argument("--plan", type=Path, help="İş Adı, Sistem, Sıklık (Günlük / Haftalık / Aylık veya saat), Azami Süre (dk)")
    p.add_argument("--an", help="Rapor anı (GG.AA.YYYY SS:DD); verilmezse şimdi")
    p.add_argument("--gun", type=int, default=7, help="İstatistik penceresi, gün (varsayılan 7)")
    p.add_argument("--tolerans", default="6", help="Gecikme toleransı, saat (varsayılan 6)")
    p.add_argument("--boyut-sapma", default="50", help="Boyut medyandan en az bu kadar %% küçükse uyarılır (varsayılan 50)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "yedekleme_dogrulama.xlsx")
    a = p.parse_args(argv)
    ornek = a.gecmis is None
    log_yol, plan_yol = (ORNEK / "gecmis.csv", a.plan or ORNEK / "plan.csv") if ornek else (a.gecmis, a.plan)
    an = zaman(a.an) if a.an else datetime(2026, 10, 12, 9, 0) if ornek else datetime.now().replace(microsecond=0)
    tol, sapma = olcu(a.tolerans), olcu(a.boyut_sapma)
    if an is None or tol is None or tol < 0 or sapma is None or not 0 < sapma < 100 or not 1 <= a.gun <= 366:
        print('[X] --an "GG.AA.YYYY SS:DD" biçiminde, --tolerans sıfır veya pozitif, --boyut-sapma 0–100 arasında, --gun 1–366 arasında olmalı')
        return 2
    for y in (log_yol, plan_yol):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(log_yol, plan_yol, a.cikti, an, a.gun, tol, sapma)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    s = r["say"]
    print(f"[OK] {len(r['isler'])} iş · {s['Sorunlu']} sorunlu, {s['Dikkat']} dikkat, {s['Tamam']} tamam · pencerede {len(r['sorunlu'])} sorunlu koşu" + ("" if plan_yol else " · plan verilmedi, gecikme denetimi yapılmadı"))
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
