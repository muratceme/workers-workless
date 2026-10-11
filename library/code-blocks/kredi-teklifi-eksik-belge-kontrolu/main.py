"""
Kredi Teklifi Eksik Belge Kontrolü — Workers / Workless kod bloğu
Krediler Tahsis › Kredi Tahsis Uzman Yardımcısı

Şubeden gelen kredi tekliflerini teklif tipine, müşteri tipine, teminata ve tutara göre zorunlu belge ve bilgi listesiyle
karşılaştırır; eksikleri şubeye iletilecek biçimde çıkarır.
  - Zorunlu belge ve bilgiler KURAL DOSYASINDAN okunur (kodda sabit liste yoktur):
      Tür (Belge / Bilgi), Ad, Teklif Tipi, Müşteri Tipi, Teminat, Asgari Tutar, Geçerlilik (gün), Açıklama
    "*" veya boş = tümü. Teminat koşulu, teklifin teminatlarından biri eşleşirse sağlanır. Asgari Tutar: teklif tutarı bu
    tutara eşit veya büyükse kural uygulanır.
  - Belge durumu: Eksik · Süresi geçmiş (belge tarihi, kontrol tarihinden "Geçerlilik (gün)" gününden eski) ·
    Tarih yok (süreli belge, tarih girilmemiş) · Tamam. Aynı belgeden birden çok kayıt varsa en yenisi esas alınır.
  - Bilgi durumu: teklif dosyasında kuraldaki "Ad" ile aynı adlı sütun boşsa Eksik; sütun hiç yoksa uyarı.
  - Sonuç: eksiği olmayan teklif "İncelemeye hazır", diğerleri "Şubeye iade".
  - Çıktı: Excel raporu + şubelere gönderilecek hazır metin (sube_iade_notlari.txt).
Kredi tahsis sürecinde istenecek belgeler bankanızın kredi politikası ve mevzuatla belirlenir; bu paket yalnız sizin
tanımladığınız listeyi tekliflerle karşılaştırır. Kredi değerlendirmesi yapmaz. İnternete bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: 12 teklif
    python main.py --teklifler t.xlsx --belgeler b.xlsx --kurallar kurallar.xlsx --tarih 09.10.2026
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


TEKLIF_SUTUNLARI = {"no": ("teklif no", "teklif numarasi", "teklif"), "sube": ("sube",), "musteri": ("musteri", "musteri adi", "unvan", "firma"),
                    "tip": ("teklif tipi", "teklif turu", "talep tipi"), "mtip": ("musteri tipi", "musteri turu"), "tutar": ("tutar", "teklif tutari", "limit", "talep edilen limit"),
                    "teminat": ("teminat", "teminat turu", "teminatlar"), "gonderen": ("gonderen", "hazirlayan", "portfoy yoneticisi")}
BELGE_SUTUNLARI = {"no": ("teklif no", "teklif numarasi", "teklif"), "belge": ("belge", "belge turu", "evrak"), "tarih": ("belge tarihi", "tarih", "duzenleme tarihi")}
KURAL_SUTUNLARI = {"tur": ("tur", "kural turu"), "ad": ("ad", "belge bilgi", "belge", "bilgi"), "tip": ("teklif tipi",), "mtip": ("musteri tipi",), "teminat": ("teminat",),
                   "asgari": ("asgari tutar", "tutar esigi"), "gun": ("gecerlilik", "gecerlilik gun", "gecerlilik suresi gun"), "aciklama": ("aciklama", "not")}
HEPSI = {"", "*", "tumu", "hepsi", "tum"}
DURUM_RENK = {"Eksik": KIRMIZI, "Süresi geçmiş": KIRMIZI, "Tarih yok": SARI, "Tamam": YESIL}


def liste(x) -> list[str]:
    return [p.strip() for p in re.split(r"[,|+]", metin(x)) if p.strip()]


def oku(t_yol: Path, b_yol: Path, k_yol: Path) -> tuple[list[dict], dict, list[dict], list[dict]]:
    uy, teklifler, belgeler, kurallar = [], [], defaultdict(list), []
    for r in kayitlar(k_yol, KURAL_SUTUNLARI, ("tur", "ad")):
        if not metin(r["ad"]):
            continue
        tur = "Bilgi" if katla(r["tur"]).startswith("bilgi") else "Belge" if katla(r["tur"]).startswith(("belge", "evrak")) else None
        gun, asgari = para(r["gun"]), para(r["asgari"])
        if tur is None or (metin(r["gun"]) and (gun is None or gun <= 0)) or (metin(r["asgari"]) and asgari is None):
            uy.append(uyari("Yüksek", "Kural okunamadı", r["ad"], f"Kural satırı {r['_satir']}: tür '{metin(r['tur'])}' (Belge / Bilgi), geçerlilik '{metin(r['gun'])}', "
                                                                 f"asgari tutar '{metin(r['asgari'])}'. Kural atlandı."))
            continue
        kurallar.append({"tur": tur, "ad": metin(r["ad"]), "tip": katla(r["tip"]), "mtip": katla(r["mtip"]), "teminat": katla(r["teminat"]), "asgari": asgari,
                         "gun": int(gun) if gun else None, "aciklama": metin(r["aciklama"])})
    if not kurallar:
        raise ValueError(f"{k_yol.name}: kural bulunamadı")
    # teklif dosyası: tanımlı sütunlar + bilgi kuralları için tüm sütunlar
    tablo = tablo_oku(t_yol)
    satirlar = kayitlar(t_yol, TEKLIF_SUTUNLARI, ("no", "tip"))
    sutun = {katla(c): j for j, c in enumerate(tablo[satirlar[0]["_satir"] - 2]) if metin(c)} if satirlar else {}      # başlık satırı
    gorulen = set()
    for r in satirlar:
        no, ham = metin(r["no"]), tablo[r["_satir"] - 1]
        if not no:
            continue
        if no in gorulen:
            uy.append(uyari("Orta", "Mükerrer teklif", no, f"Satır {r['_satir']}: teklif numarası tekrar ediyor; ilk satır kullanıldı."))
            continue
        gorulen.add(no)
        tutar = para(r["tutar"])
        if metin(r["tutar"]) and tutar is None:
            uy.append(uyari("Orta", "Tutar okunamadı", no, f"'{metin(r['tutar'])}'; tutara bağlı kurallar uygulanmadı."))
        teklifler.append({"no": no, "sube": metin(r["sube"]) or "Belirtilmemiş", "musteri": metin(r["musteri"]), "tip": metin(r["tip"]), "mtip": metin(r["mtip"]), "tutar": tutar,
                          "teminat": liste(r["teminat"]), "gonderen": metin(r["gonderen"]), "alanlar": {k: (ham[j] if j < len(ham) else None) for k, j in sutun.items()}})
    for r in kayitlar(b_yol, BELGE_SUTUNLARI, ("no", "belge")):
        no = metin(r["no"])
        if not no or not metin(r["belge"]):
            continue
        if no not in gorulen:
            uy.append(uyari("Bilgi", "Teklif listesinde yok", no, f"Belge satırı {r['_satir']}: '{metin(r['belge'])}' listede olmayan bir teklife ait."))
            continue
        t = tarih(r["tarih"])
        if metin(r["tarih"]) and not t:
            uy.append(uyari("Orta", "Tarih okunamadı", no, f"Belge satırı {r['_satir']}: {metin(r['belge'])} tarihi '{metin(r['tarih'])}'."))
        belgeler[no].append({"belge": metin(r["belge"]), "tarih": t})
    eksik_sutun = sorted({k["ad"] for k in kurallar if k["tur"] == "Bilgi" and katla(k["ad"]) not in sutun})
    for ad in eksik_sutun:
        uy.append(uyari("Yüksek", "Bilgi sütunu teklif dosyasında yok", ad, "Kural dosyasında bu bilgi zorunlu, ancak teklif dosyasında aynı adlı sütun yok; tüm tekliflerde eksik sayıldı."))
    return teklifler, belgeler, kurallar, uy


def uygulanir(k: dict, t: dict) -> bool:
    return ((k["tip"] in HEPSI or k["tip"] == katla(t["tip"])) and (k["mtip"] in HEPSI or k["mtip"] == katla(t["mtip"]))
            and (k["teminat"] in HEPSI or k["teminat"] in {katla(x) for x in t["teminat"]}) and (k["asgari"] is None or (t["tutar"] is not None and t["tutar"] >= k["asgari"])))


def neden(k: dict) -> str:
    p = [x for x in (k["tip"] not in HEPSI and "teklif tipi", k["mtip"] not in HEPSI and "müşteri tipi", k["teminat"] not in HEPSI and "teminat",
                     k["asgari"] is not None and f"tutar ≥ {k['asgari']:,.0f}".replace(",", ".")) if x]
    return ", ".join(p) if p else "tüm teklifler"


def kontrol(teklifler: list[dict], belgeler: dict, kurallar: list[dict], bugun: date) -> dict:
    uy, satirlar = [], []
    tanimli_tip = {k["tip"] for k in kurallar} - HEPSI
    for t in teklifler:
        if tanimli_tip and katla(t["tip"]) not in tanimli_tip:
            uy.append(uyari("Orta", "Teklif tipi için kural yok", t["no"], f"'{t['tip']}' tipine özel kural yok; yalnız genel kurallar uygulandı (yazım farkı olabilir)."))
        dosya = defaultdict(list)
        for b in belgeler.get(t["no"], []):
            dosya[katla(b["belge"])].append(b)
        gereken: dict = {}
        for k in kurallar:
            if uygulanir(k, t):
                x = gereken.setdefault((k["tur"], katla(k["ad"])), dict(k) | {"nedenler": []})
                if k["gun"] and (x["gun"] is None or k["gun"] < x["gun"]):
                    x["gun"] = k["gun"]                                    # birden çok kural: en kısa geçerlilik
                x["nedenler"].append(neden(k))
        for (tur, anahtar), k in gereken.items():
            son, yas = None, None
            if tur == "Bilgi":
                durum = "Tamam" if metin(t["alanlar"].get(anahtar)) else "Eksik"
            else:
                kayit = dosya.get(anahtar, [])
                son = max((b["tarih"] for b in kayit if b["tarih"]), default=None)
                yas = (bugun - son).days if son else None
                durum = "Eksik" if not kayit else "Tamam" if not k["gun"] else "Tarih yok" if son is None else "Süresi geçmiş" if yas > k["gun"] else "Tamam"
            satirlar.append({"t": t, "tur": tur, "ad": k["ad"], "durum": durum, "tarih": son, "yas": yas, "gun": k["gun"], "neden": "; ".join(dict.fromkeys(k["nedenler"])),
                             "aciklama": k["aciklama"]})
        for anahtar, kayit in dosya.items():
            if ("Belge", anahtar) not in gereken:
                uy.append(uyari("Bilgi", "İstenmeyen belge", t["no"], f"Dosyada '{kayit[0]['belge']}' var; bu teklif için zorunlu belgeler arasında değil (ad farkı olabilir)."))
    ozet = []
    for t in teklifler:
        l = [s for s in satirlar if s["t"] is t]
        sorunlu = [s for s in l if s["durum"] != "Tamam"]
        ozet.append({"t": t, "gereken": len(l), "eksik": len(sorunlu), "belge_eksik": sum(s["tur"] == "Belge" for s in sorunlu), "bilgi_eksik": sum(s["tur"] == "Bilgi" for s in sorunlu),
                     "durum": "Kural yok" if not l else "Şubeye iade" if sorunlu else "İncelemeye hazır", "sorunlu": sorunlu})
    return {"satirlar": satirlar, "ozet": ozet, "uyarilar": uy, "bugun": bugun}


def iade_metni(r: dict) -> str:
    """Şubelere gönderilecek düz metin: şube → teklif → eksikler."""
    bol = []
    for sube in sorted({o["t"]["sube"] for o in r["ozet"] if o["durum"] == "Şubeye iade"}):
        bol.append(f"=== {sube} Şubesi ===")
        for o in (o for o in r["ozet"] if o["t"]["sube"] == sube and o["durum"] == "Şubeye iade"):
            t = o["t"]
            bol.append(f"\nTeklif No: {t['no']} · {t['musteri']} · {t['tip']}" + (f" · {t['tutar']:,.0f} TL".replace(",", ".") if t["tutar"] is not None else ""))
            bol.append("Teklif, aşağıdaki eksikler tamamlanmadan incelemeye alınamamaktadır:")
            for n, s in enumerate(o["sorunlu"], 1):
                if s["tur"] == "Bilgi":
                    ek = "teklif formunda doldurulmamış"
                elif s["durum"] == "Eksik":
                    ek = "belge gönderilmemiş"
                elif s["durum"] == "Tarih yok":
                    ek = f"belge tarihi belirtilmemiş (en çok {s['gun']} gün eski olabilir)"
                else:
                    ek = f"belge {s['tarih']:%d.%m.%Y} tarihli, {s['yas']} gün eski (en çok {s['gun']} gün olabilir); güncelini gönderin"
                bol.append(f"  {n}. {s['ad']}: {ek}." + (f" ({s['aciklama']})" if s["aciklama"] else ""))
        bol.append("")
    if not bol:
        return "İade edilecek teklif yok.\n"
    return f"Kredi teklifi eksik listesi · kontrol tarihi {r['bugun']:%d.%m.%Y}\n\n" + "\n".join(bol)


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    oz = wb.active
    oz.title = "Teklif Özeti"
    _baslik(oz, ["Teklif No", "Şube", "Müşteri", "Teklif Tipi", "Müşteri Tipi", "Tutar", "Teminat", "Gereken", "Eksik Belge", "Eksik Bilgi", "Sonuç", "Gönderen"],
            (14, 14, 30, 16, 13, 18, 26, 8, 8, 8, 18, 18))
    for o in sorted(r["ozet"], key=lambda o: (-o["eksik"], o["t"]["no"])):
        t = o["t"]
        _satir(oz, [t["no"], t["sube"], t["musteri"], t["tip"], t["mtip"], t["tutar"], ", ".join(t["teminat"]), o["gereken"], o["belge_eksik"], o["bilgi_eksik"], o["durum"], t["gonderen"]],
               para_s=(6,), dolgu={11: YESIL if o["durum"] == "İncelemeye hazır" else KIRMIZI})
    oz.auto_filter.ref = f"A1:L{oz.max_row}"

    ek = wb.create_sheet("Eksikler")
    _baslik(ek, ["Teklif No", "Şube", "Müşteri", "Tür", "Belge / Bilgi", "Durum", "Belge Tarihi", "Yaş (gün)", "Geçerlilik (gün)", "Neden İsteniyor", "Açıklama", "Şubeye Bildirildi",
                 "Tamamlandı"], (14, 14, 30, 8, 34, 14, 11, 8, 10, 30, 40, 11, 11))
    for s in r["satirlar"]:
        if s["durum"] == "Tamam":
            continue
        t = s["t"]
        _satir(ek, [t["no"], t["sube"], t["musteri"], s["tur"], s["ad"], s["durum"], s["tarih"], s["yas"], s["gun"], s["neden"], s["aciklama"], "", ""], tarih_s=(7,),
               dolgu={6: DURUM_RENK[s["durum"]], 12: SARI, 13: SARI}, sar=(11,))
    ek.auto_filter.ref = f"A1:M{ek.max_row}"

    so = wb.create_sheet("Şube Özeti")
    _baslik(so, ["Şube", "Teklif", "İncelemeye Hazır", "İade", "İade Oranı %", "Toplam Eksik", "En Sık Eksik"], (16, 8, 10, 8, 10, 10, 50))
    for sube in sorted({o["t"]["sube"] for o in r["ozet"]}):
        l = [o for o in r["ozet"] if o["t"]["sube"] == sube]
        say = defaultdict(int)
        for o in l:
            for s in o["sorunlu"]:
                say[s["ad"]] += 1
        iade = sum(o["durum"] == "Şubeye iade" for o in l)
        en = sorted(say.items(), key=lambda x: (-x[1], x[0]))[:3]
        _satir(so, [sube, len(l), sum(o["durum"] == "İncelemeye hazır" for o in l), iade, yuvarla(Decimal(iade) / len(l) * 100, "0.1"), sum(o["eksik"] for o in l),
                    ", ".join(f"{a} ({n})" for a, n in en)], sar=(7,))
    so.append([])
    _baslik(so, ["En sık eksikler (tüm şubeler)", "Teklif Sayısı"], dondur=None)
    genel = defaultdict(int)
    for s in r["satirlar"]:
        if s["durum"] != "Tamam":
            genel[f"{s['ad']} ({s['tur']})"] += 1
    for a, n in sorted(genel.items(), key=lambda x: (-x[1], x[0]))[:10]:
        _satir(so, [a, n])
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)
    cikti.with_name("sube_iade_notlari.txt").write_text(iade_metni(r), encoding="utf-8")


def calistir(t_yol: Path, b_yol: Path, k_yol: Path, cikti: Path, bugun: date) -> dict:
    teklifler, belgeler, kurallar, uy = oku(t_yol, b_yol, k_yol)
    if not teklifler:
        raise ValueError(f"{t_yol.name}: teklif bulunamadı")
    r = kontrol(teklifler, belgeler, kurallar, bugun)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Kredi tekliflerini zorunlu belge ve bilgi listesiyle karşılaştırır; eksikleri şubeye iletilecek biçimde çıkarır.")
    p.add_argument("--teklifler", type=Path, default=ORNEK / "teklifler.csv", help="Teklif No, Şube, Müşteri, Teklif Tipi, Müşteri Tipi, Tutar, Teminat + zorunlu bilgi sütunları")
    p.add_argument("--belgeler", type=Path, default=ORNEK / "belgeler.csv", help="Teklif No, Belge, Belge Tarihi")
    p.add_argument("--kurallar", type=Path, default=ORNEK / "kurallar.csv", help="Tür (Belge / Bilgi), Ad, Teklif Tipi, Müşteri Tipi, Teminat, Asgari Tutar, Geçerlilik (gün), Açıklama")
    p.add_argument("--tarih", help="Kontrol tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 09.10.2026)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "teklif_eksik_kontrolu.xlsx")
    a = p.parse_args(argv)
    for y in (a.teklifler, a.belgeler, a.kurallar):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    bugun = tarih(a.tarih) if a.tarih else (date(2026, 10, 9) if a.teklifler == ORNEK / "teklifler.csv" else date.today())
    if not bugun:
        print("[X] --tarih GG.AA.YYYY biçiminde olmalı")
        return 2
    try:
        r = calistir(a.teklifler, a.belgeler, a.kurallar, a.cikti, bugun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    iade = [o for o in r["ozet"] if o["durum"] == "Şubeye iade"]
    print(f"[OK] {len(r['ozet'])} teklif · incelemeye hazır {sum(o['durum'] == 'İncelemeye hazır' for o in r['ozet'])} · şubeye iade {len(iade)}")
    for o in iade:
        print(f"     {o['t']['no']} ({o['t']['sube']}): {o['belge_eksik']} belge, {o['bilgi_eksik']} bilgi eksik")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['kim']} · {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[OK] Şube notları: {a.cikti.with_name('sube_iade_notlari.txt').resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
