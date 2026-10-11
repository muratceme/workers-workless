import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 9)


def kural(tur, ad, tip="*", mtip="*", teminat="*", asgari=None, gun=None):
    return {"tur": tur, "ad": ad, "tip": main.katla(tip), "mtip": main.katla(mtip), "teminat": main.katla(teminat), "asgari": D(asgari) if asgari else None, "gun": gun, "aciklama": ""}


KURALLAR = [kural("Belge", "Talep Formu"), kural("Belge", "Mali Tablolar", mtip="Tüzel", gun=150), kural("Belge", "Ekspertiz Raporu", teminat="İpotek", gun=180),
            kural("Belge", "Denetim Raporu", asgari=50_000_000), kural("Belge", "Limit Dökümü", tip="Limit Artışı"), kural("Bilgi", "Kullanım Amacı"),
            kural("Bilgi", "Risk Grubu", mtip="Tüzel", asgari=10_000_000), kural("Belge", "Mali Tablolar", tip="Yenileme", gun=90)]


def teklif(no="T1", tip="Yeni Limit", mtip="Tüzel", tutar=5_000_000, teminat=("Kefalet",), sube="Merkez", **alanlar):
    return {"no": no, "sube": sube, "musteri": "M", "tip": tip, "mtip": mtip, "tutar": D(tutar) if tutar is not None else None, "teminat": list(teminat), "gonderen": "",
            "alanlar": {main.katla(k): v for k, v in alanlar.items()}}


def belge(ad, gun_once=None):
    return {"belge": ad, "tarih": date.fromordinal(BUGUN.toordinal() - gun_once) if gun_once is not None else None}


def durumlar(r):
    return {s["ad"]: s["durum"] for s in r["satirlar"]}


class KuralTesti(unittest.TestCase):
    def test_kosullar(self):
        def gereken(t):
            return sorted(k["ad"] for k in KURALLAR if main.uygulanir(k, t))
        self.assertEqual(gereken(teklif()), ["Kullanım Amacı", "Mali Tablolar", "Talep Formu"])
        self.assertEqual(gereken(teklif(mtip="Şahıs")), ["Kullanım Amacı", "Talep Formu"])
        self.assertEqual(gereken(teklif(tutar=50_000_000, teminat=("İPOTEK", "Kefalet"))),
                         ["Denetim Raporu", "Ekspertiz Raporu", "Kullanım Amacı", "Mali Tablolar", "Risk Grubu", "Talep Formu"])
        self.assertEqual(gereken(teklif(tip="limit artışı", tutar=9_999_999)), ["Kullanım Amacı", "Limit Dökümü", "Mali Tablolar", "Talep Formu"])
        self.assertEqual(gereken(teklif(tutar=None)), ["Kullanım Amacı", "Mali Tablolar", "Talep Formu"])       # tutar yoksa eşikli kural uygulanmaz

    def test_belge_durumlari(self):
        t = teklif(teminat=("İpotek",), **{"Kullanım Amacı": "Stok"})
        b = {"T1": [belge("TALEP FORMU"), belge("Mali Tablolar", 150), belge("Ekspertiz Raporu", 181)]}
        self.assertEqual(durumlar(main.kontrol([t], b, KURALLAR, BUGUN)), {"Talep Formu": "Tamam", "Mali Tablolar": "Tamam", "Ekspertiz Raporu": "Süresi geçmiş", "Kullanım Amacı": "Tamam"})
        b = {"T1": [belge("Mali Tablolar"), belge("Ekspertiz Raporu", 400), belge("Ekspertiz Raporu", 10)]}
        r = main.kontrol([t], b, KURALLAR, BUGUN)
        self.assertEqual(durumlar(r), {"Talep Formu": "Eksik", "Mali Tablolar": "Tarih yok", "Ekspertiz Raporu": "Tamam", "Kullanım Amacı": "Tamam"})   # en yeni kayıt
        self.assertEqual((r["ozet"][0]["durum"], r["ozet"][0]["belge_eksik"], r["ozet"][0]["bilgi_eksik"]), ("Şubeye iade", 2, 0))

    def test_en_kisa_gecerlilik(self):
        t = teklif(tip="Yenileme", **{"Kullanım Amacı": "x"})
        r = main.kontrol([t], {"T1": [belge("Talep Formu"), belge("Mali Tablolar", 100)]}, KURALLAR, BUGUN)
        s = next(x for x in r["satirlar"] if x["ad"] == "Mali Tablolar")
        self.assertEqual((s["durum"], s["gun"], s["yas"]), ("Süresi geçmiş", 90, 100))                           # 150 ve 90 günlük iki kural → 90

    def test_bilgi_ve_sonuc(self):
        t1 = teklif("T1", mtip="Şahıs", **{"Kullanım Amacı": "  "})
        t2 = teklif("T2", mtip="Şahıs", **{"Kullanım Amacı": "Stok"})
        t3 = teklif("T3", tip="Bilinmeyen", mtip="Şahıs", **{"Kullanım Amacı": "Stok"})
        r = main.kontrol([t1, t2, t3], {x: [belge("Talep Formu")] for x in ("T1", "T2", "T3")} | {"T2": [belge("Talep Formu"), belge("Pasaport")]}, KURALLAR, BUGUN)
        self.assertEqual([o["durum"] for o in r["ozet"]], ["Şubeye iade", "İncelemeye hazır", "İncelemeye hazır"])
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]), [("Teklif tipi için kural yok", "T1"), ("Teklif tipi için kural yok", "T2"), ("Teklif tipi için kural yok", "T3"), ("İstenmeyen belge", "T2")])   # test kurallarında 'Yeni Limit'e özel kural yok

    def test_iade_metni(self):
        t = teklif(teminat=("İpotek",), sube="Sanayi", **{"Kullanım Amacı": ""})
        r = main.kontrol([t, teklif("T2", mtip="Şahıs", **{"Kullanım Amacı": "x"})], {"T1": [belge("Mali Tablolar"), belge("Ekspertiz Raporu", 200)], "T2": [belge("Talep Formu")]},
                         KURALLAR, BUGUN)
        m = main.iade_metni(r)
        self.assertIn("=== Sanayi Şubesi ===", m)
        self.assertIn("Teklif No: T1 · M · Yeni Limit · 5.000.000 TL", m)
        self.assertIn("1. Talep Formu: belge gönderilmemiş.", m)
        self.assertIn("2. Mali Tablolar: belge tarihi belirtilmemiş (en çok 150 gün eski olabilir).", m)
        self.assertIn("3. Ekspertiz Raporu: belge 23.03.2026 tarihli, 200 gün eski (en çok 180 gün olabilir); güncelini gönderin.", m)
        self.assertIn("4. Kullanım Amacı: teklif formunda doldurulmamış.", m)
        self.assertNotIn("T2", m)
        self.assertEqual(main.iade_metni(main.kontrol([], {}, KURALLAR, BUGUN)), "İade edilecek teklif yok.\n")


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "teklifler.csv", ORNEK / "belgeler.csv", ORNEK / "kurallar.csv", cikti, BUGUN)
            d = {o["t"]["no"]: o["durum"] for o in r["ozet"]}
            self.assertEqual((len(d), sum(v == "İncelemeye hazır" for v in d.values()), sum(v == "Şubeye iade" for v in d.values())), (12, 5, 7))
            self.assertTrue(all(d[x] == "İncelemeye hazır" for x in ("TK-2602", "TK-2606", "TK-2607", "TK-2611")))
            buyuk = [s["ad"] for s in r["satirlar"] if s["t"]["no"] == "TK-2611"]
            self.assertIn("Bağımsız Denetim Raporu", buyuk)                           # 90 milyon ≥ 50 milyon
            self.assertNotIn("Bağımsız Denetim Raporu", [s["ad"] for s in r["satirlar"] if s["t"]["no"] == "TK-2603"])
            self.assertEqual(sorted(u["tur"] for u in r["uyarilar"]), ["Teklif listesinde yok", "Teklif tipi için kural yok", "İstenmeyen belge"])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Teklif Özeti", "Eksikler", "Şube Özeti", "Uyarılar"])
            self.assertIn("TK-2601", (Path(tmp) / "sube_iade_notlari.txt").read_text(encoding="utf-8"))

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            t, b, k = (Path(tmp) / x for x in ("t.csv", "b.csv", "k.csv"))
            t.write_text("Teklif No;Şube;Teklif Tipi;Müşteri Tipi;Tutar;Teminat;Kullanım Amacı\nA1;Merkez;Yeni Limit;Tüzel;1.500.000,00;İpotek, Kefalet;Stok\n"
                         "A1;Merkez;Yeni Limit;Tüzel;1;;\nA2;Merkez;Yeni Limit;Şahıs;çok;;\n", encoding="utf-8")
            b.write_text("Teklif No;Belge;Belge Tarihi\nA1;Talep Formu;dün\nA9;Talep Formu;01.10.2026\n", encoding="utf-8")
            k.write_text("Tür;Ad;Teklif Tipi;Müşteri Tipi;Teminat;Asgari Tutar;Geçerlilik (gün)\nBelge;Talep Formu;*;*;*;;\nBilgi;Kullanım Amacı;;;;;\nBilgi;Vade;*;*;*;;\n"
                         "Rapor;Yanlış Tür;*;*;*;;\nBelge;Bozuk;*;*;*;;otuz\n", encoding="utf-8")
            tk, bl, kr, uy = main.oku(t, b, k)
            self.assertEqual([(x["no"], x["tutar"], x["teminat"], x["alanlar"]["kullanim amaci"]) for x in tk], [("A1", D("1500000.00"), ["İpotek", "Kefalet"], "Stok"), ("A2", None, [], "")])
            self.assertEqual(len(kr), 3)
            self.assertEqual(sorted(u["tur"] for u in uy), ["Bilgi sütunu teklif dosyasında yok", "Kural okunamadı", "Kural okunamadı", "Mükerrer teklif", "Tarih okunamadı",
                                                            "Teklif listesinde yok", "Tutar okunamadı"])
            arg = ["--teklifler", str(t), "--belgeler", str(b), "--kurallar", str(k), "--cikti", str(Path(tmp) / "x.xlsx")]
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(arg + ["--tarih", "bugün"]), 2)
            self.assertEqual(main.main(["--kurallar", str(Path(tmp) / "yok.csv")]), 1)
            k.write_text("Tür;Ad\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "o.xlsx")]), 0)


if __name__ == "__main__":
    unittest.main()
