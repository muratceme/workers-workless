import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)


def s(kod, cikis=None, giris=None, miktar=10, maliyet=5, depo="D", kategori="K"):
    """cikis / giris: kaç gün önce (None = tarih yok)."""
    g = lambda n: None if n is None else BUGUN - timedelta(days=n)   # noqa: E731
    return {"kod": kod, "ad": "Ad", "depo": depo, "kategori": kategori, "miktar": D(str(miktar)), "maliyet": None if maliyet is None else D(str(maliyet)), "cikis": g(cikis), "giris": g(giris)}


class YaslandirTesti(unittest.TestCase):
    def test_dilim_ve_durum(self):
        gunler = (0, 30, 31, 60, 61, 89, 90, 179, 180, 365, 366, -5)
        r = main.yaslandir([s(str(g), g) for g in gunler], BUGUN)
        x = {v["kod"]: (v["yas"], v["dilim"], v["durum"]) for v in r["stoklar"]}
        self.assertEqual([x[str(g)] for g in gunler],
                         [(0, "0–30", "Normal"), (30, "0–30", "Normal"), (31, "31–60", "Normal"), (60, "31–60", "Normal"), (61, "61–90", "Normal"), (89, "61–90", "Normal"), (90, "61–90", "Yavaş"), (179, "91–180", "Yavaş"),
                          (180, "91–180", "Hareketsiz"), (365, "181–365", "Hareketsiz"), (366, "366+", "Hareketsiz"), (0, "0–30", "Normal")])
        self.assertEqual([(d["ad"], d["kalem"], d["tutar"], d["pay"]) for d in r["dilimler"]],
                         [("0–30", 3, D(150), D("25.00")), ("31–60", 2, D(100), D("16.67")), ("61–90", 3, D(150), D("25.00")), ("91–180", 2, D(100), D("16.67")), ("181–365", 1, D(50), D("8.33")), ("366+", 1, D(50), D("8.33"))])
        self.assertEqual([(d["ad"], d["kalem"], d["tutar"]) for d in r["durumlar"]], [("Hareketsiz", 3, D(150)), ("Yavaş", 2, D(100)), ("Normal", 7, D(350)), ("Tarih yok", 0, D(0))])
        self.assertEqual(([v["kod"] for v in r["stoklar"]][:3], [(u["tur"], u["kim"]) for u in r["uyarilar"]], r["toplam"]), (["366", "365", "180"], [("Hareket tarihi ileride", "-5 · Ad")], D(600)))
        r = main.yaslandir([s("a", 45), s("b", 46), s("c", 119), s("d", 120)], BUGUN, dilimler=(45, 120), yavas=46, hareketsiz=120)
        self.assertEqual(sorted((v["kod"], v["dilim"], v["durum"]) for v in r["stoklar"]), [("a", "0–45", "Normal"), ("b", "46–120", "Yavaş"), ("c", "46–120", "Yavaş"), ("d", "46–120", "Hareketsiz")])
        self.assertEqual([d["ad"] for d in r["dilimler"]], ["0–45", "46–120", "121+"])

    def test_tarih_ve_miktar(self):
        l = [s("girişli", None, 100), s("yeni_giriş", None, 20), s("tarihsiz"), s("sıfır", 5, miktar=0), s("eksi", 5, miktar=-1), s("maliyetsiz", 5, maliyet=None), s("ondalık", 5, miktar="2.5", maliyet="3.333"), s("ikisi", 200, 10)]
        r = main.yaslandir(l, BUGUN)
        x = {v["kod"]: (v["yas"], v["dilim"], v["durum"], v["tutar"]) for v in r["stoklar"]}
        self.assertEqual(x, {"girişli": (100, "91–180", "Yavaş", D("50.00")), "yeni_giriş": (20, "0–30", "Normal", D("50.00")), "tarihsiz": (None, "(tarih yok)", "Tarih yok", D("50.00")), "maliyetsiz": (5, "0–30", "Normal", D("0.00")),
                             "ondalık": (5, "0–30", "Normal", D("8.33")), "ikisi": (200, "181–365", "Hareketsiz", D("50.00"))})
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         sorted([("Orta", "Hiç çıkışı olmayan stok", "girişli · Ad"), ("Orta", "Tarih yok", "tarihsiz · Ad"), ("Orta", "Eksi stok", "eksi · Ad"), ("Bilgi", "Maliyet yok", "maliyetsiz · Ad")]))
        self.assertEqual(([v["kod"] for v in r["stoklar"]][-1], r["dilimler"][-1], [v["kod"] for v in r["hareketsiz"]]), ("tarihsiz", {"ad": "(tarih yok)", "kalem": 1, "tutar": D("50.00"), "pay": D("24.00")}, ["ikisi"]))
        r = main.yaslandir([s("ikisi", 200, 10), s("yalnız_çıkış", 200)], BUGUN, esas="hareket")
        self.assertEqual(sorted((v["kod"], v["yas"], v["durum"]) for v in r["stoklar"]), [("ikisi", 10, "Normal"), ("yalnız_çıkış", 200, "Hareketsiz")])
        self.assertEqual(main.yaslandir([s("a", 5, maliyet=0)], BUGUN)["durumlar"][2]["pay"], None)
        for bos in ([], [s("a", 5, miktar=0)], [s("a", 5, miktar=-3)]):
            with self.assertRaises(ValueError):
                main.yaslandir(bos, BUGUN)

    def test_kirilim(self):
        l = [s("1", 5, depo="A", kategori="X", maliyet=10), s("2", 100, depo="a", kategori="Y", maliyet=20), s("3", 200, depo="B", kategori="X", maliyet=30), s("4", 400, depo="B", kategori="x", maliyet=40)]
        r = main.yaslandir(l, BUGUN)
        k = lambda alan: [(g["ad"], g["kalem"], g["tutar"], g["yavas"], g["hareketsiz"], g["hareketsiz_pay"]) for g in r[alan]]   # noqa: E731
        self.assertEqual((k("depolar"), k("kategoriler")), ([("A", 2, D(300), D(200), D(0), D("0.00")), ("B", 2, D(700), D(0), D(700), D("100.00"))],
                                                            [("X", 3, D(800), D(0), D(700), D("87.50")), ("Y", 1, D(200), D(200), D(0), D("0.00"))]))
        self.assertEqual([v["kod"] for v in r["hareketsiz"]], ["4", "3"])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "stok.csv", cikti, BUGUN)
            self.assertEqual((r["toplam"], len(r["stoklar"])), (D(564000), 10))
            self.assertEqual([(v["kod"], v["yas"], v["dilim"], v["durum"], v["tutar"]) for v in r["stoklar"]],
                             [("S-006", 418, "366+", "Hareketsiz", D(72000)), ("S-008", 334, "181–365", "Hareketsiz", D(128000)), ("S-005", 239, "181–365", "Hareketsiz", D(90000)), ("S-004", 124, "91–180", "Yavaş", D(52000)),
                              ("S-011", 94, "91–180", "Yavaş", D(0)), ("S-003", 72, "61–90", "Normal", D(114000)), ("S-002", 45, "31–60", "Normal", D(18000)), ("S-007", 12, "0–30", "Normal", D(45000)), ("S-001", 7, "0–30", "Normal", D(34000)),
                              ("S-009", None, "(tarih yok)", "Tarih yok", D(11000))])
            self.assertEqual([(d["ad"], d["kalem"], d["tutar"], d["pay"]) for d in r["dilimler"]],
                             [("0–30", 2, D(79000), D("14.01")), ("31–60", 1, D(18000), D("3.19")), ("61–90", 1, D(114000), D("20.21")), ("91–180", 2, D(52000), D("9.22")), ("181–365", 2, D(218000), D("38.65")),
                              ("366+", 1, D(72000), D("12.77")), ("(tarih yok)", 1, D(11000), D("1.95"))])
            self.assertEqual([(d["ad"], d["kalem"], d["tutar"], d["pay"]) for d in r["durumlar"]],
                             [("Hareketsiz", 3, D(290000), D("51.42")), ("Yavaş", 2, D(52000), D("9.22")), ("Normal", 4, D(211000), D("37.41")), ("Tarih yok", 1, D(11000), D("1.95"))])
            k = lambda alan: [(g["ad"], g["kalem"], g["tutar"], g["yavas"], g["hareketsiz"], g["hareketsiz_pay"]) for g in r[alan]]   # noqa: E731
            self.assertEqual(k("depolar"), [("Ana Depo", 6, D(308000), D(52000), D(90000), D("29.22")), ("Şube Depo", 4, D(256000), D(0), D(200000), D("78.13"))])
            self.assertEqual(k("kategoriler"), [("Ambalaj", 2, D(56000), D(0), D(0), D("0.00")), ("Mamul", 2, D(200000), D(0), D(200000), D("100.00")), ("Sarf", 3, D(166000), D(52000), D(0), D("0.00")),
                                                ("Yedek Parça", 3, D(142000), D(0), D(90000), D("63.38"))])
            self.assertEqual([v["kod"] for v in r["hareketsiz"]], ["S-008", "S-005", "S-006"])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Orta", "Eksi stok", "S-010 · Conta Seti"), ("Orta", "Hiç çıkışı olmayan stok", "S-008 · Özel Sipariş Flanş"), ("Orta", "Tarih yok", "S-009 · Etiket Rulosu"), ("Bilgi", "Maliyet yok", "S-011 · Boya Tineri")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Depolar", "Kategoriler", "Hareketsiz Stok", "Stoklar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "s.csv"
            y.write_text("Malzeme Kodu;Malzeme;Ambar;Grup;Stok Miktarı;Maliyet;Son Satış Tarihi;Son Alış Tarihi\nM1;Vida;;;1.250,5;2,40;05.10.2026;x\nM2;Somun;A;G;12.5;;;01.09.2026\nM3;X;;;x;1;;\nM4;X;;;1;-1;;\n;X;;;1;1;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["kod"], x["depo"], x["kategori"], x["miktar"], x["maliyet"], x["cikis"], x["giris"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("M1", "(depo yok)", "(kategori yok)", D("1250.5"), D("2.40"), date(2026, 10, 5), None), ("M2", "A", "G", D("12.5"), None, None, date(2026, 9, 1))],
                              [("Tarih okunamadı", "M1 · Vida"), ("Stok satırı okunamadı", "M3 · X"), ("Stok satırı okunamadı", "M4 · X")]))
            self.assertEqual([main.dilim_coz(x) for x in ("30;60", "90", "60;30", "x", "")], [(30, 60), (90,), None, None, None])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--stok", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "12.10.2026", "--dilimler", "30;90;180;365", "--yavas", "60", "--hareketsiz", "120", "--esas", "hareket"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + x) for x in (["--tarih", "x"], ["--dilimler", "90;30"], ["--yavas", "0"], ["--yavas", "180"], ["--hareketsiz", "90"])], [2] * 5)
            self.assertEqual(main.main(["--stok", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Stok Kodu;Depo\nA;1\n", "Stok Kodu;Miktar\n;1\n", "Stok Kodu;Miktar\nA;0\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
