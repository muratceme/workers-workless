"""
Stok Yaşlandırma Raporu — Gerisi Bende kod bloğu
Lojistik ve Depo › Depo Sorumlusu

Son hareket tarihine göre hareketsiz ve yavaş dönen stokları tutar bazında yaşlandırır.
  Stok tutarı   = miktar × birim maliyet
  Yaş (gün)     = durum tarihi − son çıkış tarihi   (son çıkış yoksa son giriş tarihi; --esas hareket: ikisinden yeni olan)
  Yaş dilimi    = 0–30, 31–60, 61–90, 91–180, 181–365, 366+ gün   (--dilimler ile değişir)
  Durum         = Hareketsiz (yaş ≥ 180) · Yavaş (yaş ≥ 90) · Normal · Tarih yok
  Pay %         = dilimin (ya da durumun) tutarı ÷ toplam stok tutarı × 100
Eşikler yönetim tercihidir. Rapor değer düşüklüğü karşılığı hesaplamaz; yalnız hangi stokun ne kadar süredir hareket görmediğini gösterir.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
SUTUN = {"kod": ("stok kodu", "malzeme kodu", "urun kodu", "kod"), "ad": ("stok adi", "malzeme", "malzeme adi", "urun", "ad"), "depo": ("depo", "ambar", "lokasyon"), "kategori": ("kategori", "grup", "malzeme grubu"),
         "miktar": ("miktar", "stok miktari", "eldeki miktar"), "maliyet": ("birim maliyet", "maliyet", "birim fiyat"), "cikis": ("son cikis tarihi", "son cikis", "son satis tarihi", "son tuketim tarihi"),
         "giris": ("son giris tarihi", "son giris", "son alis tarihi")}
DURUMLAR = ("Hareketsiz", "Yavaş", "Normal", "Tarih yok")
DURUM_RENK = {"Hareketsiz": KIRMIZI, "Yavaş": SARI, "Normal": None, "Tarih yok": MAVI}
Z, YUZ = Decimal(0), Decimal(100)


def olcu(x) -> Decimal | None:
    """Miktar okur: virgül ya da tek nokta ondalık ayırıcıdır; '1.250,5' binlik ayırıcılıdır."""
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x or ""))
    try:
        return Decimal(s.replace(".", "").replace(",", ".") if "," in s else s) if s else None
    except InvalidOperation:
        return None


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("kod", "miktar")):
        kod = metin(r["kod"])
        if not kod:
            continue
        kim = f"{kod} · {metin(r['ad'])}".strip(" ·")
        miktar, maliyet = olcu(r["miktar"]), para(r["maliyet"])
        if miktar is None or (metin(r["maliyet"]) and (maliyet is None or maliyet < 0)):
            uy.append(uyari("Yüksek", "Stok satırı okunamadı", kim, f"Satır {r['_satir']}: miktar '{metin(r['miktar'])}', birim maliyet '{metin(r['maliyet'])}'. Satır alınmadı."))
            continue
        for alan, ad in (("cikis", "Son çıkış"), ("giris", "Son giriş")):
            if metin(r[alan]) and tarih(r[alan]) is None:
                uy.append(uyari("Orta", "Tarih okunamadı", kim, f"{ad} tarihi '{metin(r[alan])}' okunamadı; boş sayıldı."))
        l.append({"kod": kod, "ad": metin(r["ad"]), "depo": metin(r["depo"]) or "(depo yok)", "kategori": metin(r["kategori"]) or "(kategori yok)", "miktar": miktar, "maliyet": maliyet, "cikis": tarih(r["cikis"]), "giris": tarih(r["giris"])})
    return l, uy


def _pay(tutar: Decimal, toplam: Decimal) -> Decimal | None:
    return yuvarla(tutar * YUZ / toplam) if toplam else None


def yaslandir(stoklar: list[dict], bugun: date, dilimler: tuple[int, ...] = (30, 60, 90, 180, 365), yavas: int = 90, hareketsiz: int = 180, esas: str = "cikis") -> dict:
    uy, liste = [], []
    etiketler = [f"{a}–{b}" for a, b in zip((0,) + tuple(d + 1 for d in dilimler), dilimler)] + [f"{dilimler[-1] + 1}+", "(tarih yok)"]
    for s in stoklar:
        kim = f"{s['kod']} · {s['ad']}".strip(" ·")
        if s["miktar"] < 0:
            uy.append(uyari("Orta", "Eksi stok", kim, f"Miktar {s['miktar']:f}. Kayıt hatası olabilir; rapora alınmadı."))
            continue
        if s["miktar"] == 0:
            continue
        s["tutar"] = yuvarla(s["miktar"] * (s["maliyet"] or Z))
        tarihler = [t for t in (s["cikis"], s["giris"]) if t]
        s["referans"] = (max(tarihler) if tarihler else None) if esas == "hareket" else s["cikis"] or s["giris"]
        if s["referans"] is None:
            s["yas"], s["dilim"], s["durum"] = None, "(tarih yok)", "Tarih yok"
            uy.append(uyari("Orta", "Tarih yok", kim, f"Son çıkış ve son giriş tarihi yazılmamış; {s['tutar']:,.2f} tutarındaki stok yaşlandırılamadı."))
        else:
            s["yas"] = max((bugun - s["referans"]).days, 0)
            s["dilim"] = etiketler[next((i for i, d in enumerate(dilimler) if s["yas"] <= d), len(dilimler))]
            s["durum"] = "Hareketsiz" if s["yas"] >= hareketsiz else "Yavaş" if s["yas"] >= yavas else "Normal"
            if s["cikis"] is None and s["yas"] >= yavas:
                uy.append(uyari("Orta", "Hiç çıkışı olmayan stok", kim, f"Son giriş {s['giris']:%d.%m.%Y} ({s['yas']} gün önce), çıkış kaydı yok. Tutar {s['tutar']:,.2f}."))
            if s["referans"] > bugun:
                uy.append(uyari("Orta", "Hareket tarihi ileride", kim, f"Son hareket tarihi {s['referans']:%d.%m.%Y}, durum tarihi {bugun:%d.%m.%Y}."))
        if not s["maliyet"]:
            uy.append(uyari("Bilgi", "Maliyet yok", kim, "Birim maliyet yazılmamış ya da sıfır; stok tutarı sıfır alındı."))
        liste.append(s)
    if not liste:
        raise ValueError("Yaşlandırılacak stok yok (miktarı sıfırdan büyük satır bulunamadı).")
    toplam = sum((s["tutar"] for s in liste), Z)

    def topla(sec, adlar) -> list[dict]:
        return [{"ad": ad, "kalem": sum(1 for s in liste if sec(s) == ad), "tutar": sum((s["tutar"] for s in liste if sec(s) == ad), Z), "pay": _pay(sum((s["tutar"] for s in liste if sec(s) == ad), Z), toplam)} for ad in adlar]

    def kirilim(alan: str) -> list[dict]:
        grup = {}
        for s in liste:
            g = grup.setdefault(katla(s[alan]), {"ad": s[alan], "kalem": 0, "tutar": Z, "yavas": Z, "hareketsiz": Z})
            g["kalem"] += 1
            g["tutar"] += s["tutar"]
            g["yavas"] += s["tutar"] if s["durum"] == "Yavaş" else Z
            g["hareketsiz"] += s["tutar"] if s["durum"] == "Hareketsiz" else Z
        for g in grup.values():
            g["hareketsiz_pay"] = _pay(g["hareketsiz"], g["tutar"])
        return [g for _, g in sorted(grup.items())]

    dilim_ozet = [d for d in topla(lambda s: s["dilim"], etiketler) if d["ad"] != "(tarih yok)" or d["kalem"]]
    depolar, kategoriler = kirilim("depo"), kirilim("kategori")   # adlar dosyadaki ilk yazımla kalsın diye sıralamadan önce
    liste.sort(key=lambda s: (s["yas"] is None, -(s["yas"] or 0), -s["tutar"], katla(s["kod"])))
    return {"stoklar": liste, "toplam": toplam, "dilimler": dilim_ozet, "durumlar": topla(lambda s: s["durum"], DURUMLAR), "depolar": depolar, "kategoriler": kategoriler,
            "hareketsiz": sorted((s for s in liste if s["durum"] == "Hareketsiz"), key=lambda s: (-s["tutar"], katla(s["kod"]))), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, bugun: date) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    _baslik(ws, ["Yaş Dilimi (gün)", "Kalem", "Stok Tutarı", "Pay %"], (22, 9, 18, 10))
    for d in r["dilimler"]:
        _satir(ws, [d["ad"], d["kalem"], d["tutar"], d["pay"]], para_s=(3, 4))
    n = _satir(ws, ["TOPLAM", len(r["stoklar"]), r["toplam"], None], para_s=(3,))
    for h in ws[n]:
        h.font = Font(bold=True)
    ws.append([])
    _baslik(ws, ["Durum", "Kalem", "Stok Tutarı", "Pay %"], dondur="A2")
    for d in r["durumlar"]:
        _satir(ws, [d["ad"], d["kalem"], d["tutar"], d["pay"]], para_s=(3, 4), dolgu={1: DURUM_RENK[d["ad"]]})
    ws.append([])
    ws.append([f"Durum tarihi: {bugun:%d.%m.%Y}"])
    for baslik, alan, ad in (("Depolar", "depolar", "Depo"), ("Kategoriler", "kategoriler", "Kategori")):
        ws = wb.create_sheet(baslik)
        _baslik(ws, [ad, "Kalem", "Stok Tutarı", "Yavaş Dönen", "Hareketsiz", "Hareketsiz Pay %"], (26, 9, 18, 16, 16, 12))
        for g in r[alan]:
            _satir(ws, [g["ad"], g["kalem"], g["tutar"], g["yavas"] or None, g["hareketsiz"] or None, g["hareketsiz_pay"]], para_s=(3, 4, 5, 6))
    basliklar = ["Stok Kodu", "Stok Adı", "Depo", "Kategori", "Miktar", "Birim Maliyet", "Stok Tutarı", "Son Çıkış", "Son Giriş", "Yaş (gün)", "Yaş Dilimi", "Durum"]
    satir = lambda s: [s["kod"], s["ad"], s["depo"], s["kategori"], s["miktar"], s["maliyet"], s["tutar"], s["cikis"], s["giris"], s["yas"], s["dilim"], s["durum"]]   # noqa: E731
    ws = wb.create_sheet("Hareketsiz Stok")
    _baslik(ws, basliklar + ["Karar (siz doldurun)"], (14, 30, 16, 16, 11, 13, 15, 12, 12, 9, 12, 12, 30))
    for s in r["hareketsiz"]:
        _satir(ws, satir(s) + [None], para_s=(5, 6, 7), tarih_s=(8, 9), dolgu={12: KIRMIZI})
    ws = wb.create_sheet("Stoklar")
    _baslik(ws, basliklar, (14, 30, 16, 16, 11, 13, 15, 12, 12, 9, 12, 12))
    for s in r["stoklar"]:
        _satir(ws, satir(s), para_s=(5, 6, 7), tarih_s=(8, 9), dolgu={12: DURUM_RENK[s["durum"]]})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, **kw) -> dict:
    stoklar, uy = oku(yol)
    r = yaslandir(stoklar, bugun, **kw)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, bugun)
    return r


def dilim_coz(s: str) -> tuple[int, ...] | None:
    try:
        d = tuple(int(x) for x in re.split(r"[;,\s]+", s.strip()) if x)
    except ValueError:
        return None
    return d if d and d[0] >= 0 and all(a < b for a, b in zip(d, d[1:])) else None


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Son hareket tarihine göre hareketsiz ve yavaş dönen stokları tutar bazında yaşlandırır.")
    p.add_argument("--stok", type=Path, help="Stok Kodu, Stok Adı, Depo, Kategori, Miktar, Birim Maliyet, Son Çıkış Tarihi, Son Giriş Tarihi")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--dilimler", default="30;60;90;180;365", help="Yaş dilimi sınırları, gün (varsayılan 30;60;90;180;365)")
    p.add_argument("--yavas", type=int, default=90, help="Bu kadar gün ve üstü hareket görmeyen stok 'Yavaş' sayılır (varsayılan 90)")
    p.add_argument("--hareketsiz", type=int, default=180, help="Bu kadar gün ve üstü hareket görmeyen stok 'Hareketsiz' sayılır (varsayılan 180)")
    p.add_argument("--esas", choices=("cikis", "hareket"), default="cikis", help="cikis: son çıkışa göre (varsayılan); hareket: son çıkış ve son girişten yeni olana göre")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "stok_yaslandirma.xlsx")
    a = p.parse_args(argv)
    ornek = a.stok is None
    yol = ORNEK / "stok.csv" if ornek else a.stok
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    dilimler = dilim_coz(a.dilimler)
    if bugun is None or dilimler is None or not 0 < a.yavas < a.hareketsiz:
        print("[X] --tarih GG.AA.YYYY; --dilimler artan tam sayılar (örn. 30;60;90); --yavas pozitif ve --hareketsiz değerinden küçük olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, bugun, dilimler=dilimler, yavas=a.yavas, hareketsiz=a.hareketsiz, esas=a.esas)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    d = {x["ad"]: x for x in r["durumlar"]}
    print(f"[OK] {len(r['stoklar'])} kalem · stok tutarı {r['toplam']:,.2f} · hareketsiz {d['Hareketsiz']['tutar']:,.2f}" + (f" (%{d['Hareketsiz']['pay']})" if d["Hareketsiz"]["pay"] is not None else "") + f" · yavaş {d['Yavaş']['tutar']:,.2f}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
