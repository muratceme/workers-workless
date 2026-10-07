"""
Banka Hareketlerinden Muhasebe Fişi Önerisi — Workers / Workless kod bloğu
Muhasebe ve Defter İşleri (Mali Müşavirlik) › Muhasebe Elemanı (Büro)

Banka ekstresindeki her hareket için karşı hesabı bulur ve çift taraflı muhasebe fişi satırları önerir:
  1. Cari eşleşmesi: açıklamadaki IBAN → VKN/TCKN → unvan (şirket türü ekleri atılarak) cari kart listesiyle
     karşılaştırılır. Giriş müşteriden tahsilat (varsayılan 120 alacak), çıkış tedarikçiye ödeme (320 borç)
     sayılır; cari kartta hesap kodu varsa o kullanılır.
  2. Kural eşleşmesi: kurallar dosyasındaki anahtar kelime, yön ve tutar aralığına göre hesap kodu
     (maaş → 335, SGK → 361, vergi → 360, banka masrafı → 780, faiz geliri → 642 ... — örnek kurallardır,
     kendi hesap planınıza göre düzenleyin).
  3. Eşleşmeyen hareketler "Manuel" olarak listelenir; fişe girmez.
Fiş: giriş → banka hesabı (102) BORÇ / karşı hesap ALACAK; çıkış → karşı hesap BORÇ / banka ALACAK.
Çıktı genel bir fiş aktarım tablosudur (Fiş No, Tarih, Hesap Kodu, Açıklama, Borç, Alacak); kendi programınızın
içe aktarma biçimine uyarlayın. İnternete bağlanmaz.

Ekstre okuma çekirdeği: Banka Mutabakatı kod bloğu (ekstre_cekirdek.py).

Kullanım:
    python main.py                                         # örnek ekstre, cari ve kurallarla dener
    python main.py --ekstre ekstre.xlsx --cariler cari_kartlar.xlsx --kurallar kurallar.csv --banka-hesap 102.01.001
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import Counter
from datetime import datetime
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import ekstre_cekirdek as ek

BURASI = Path(__file__).resolve().parent
EKLER = r"\b(a\.?\s?ş\.?|anonim|şirketi|şti\.?|ltd\.?|limited|san\.?|sanayi|tic\.?|ticaret|ve|ins\.?|inşaat|paz\.?|pazarlama|ith\.?|ihr\.?|ithalat|ihracat|gıda|tekstil|lojistik|koll\.?|kolektif)\b"


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def unvan_anahtari(s: str) -> list[str]:
    """Şirket türü ve yaygın sektör eklerini atıp ayırt edici kelimeleri döndürür."""
    k = re.sub(EKLER, " ", kucuk(s))
    return [w for w in re.findall(r"[a-zçğıöşü0-9]+", k) if len(w) >= 3]


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def carileri_oku(yol: Path | None) -> list[dict]:
    if not yol:
        return []
    s = tablo_oku(yol)
    b = [kucuk(x) for x in s[0]]
    bul = lambda *a: next((i for i, x in enumerate(b) if x in a), None)  # noqa: E731
    k = {"kod": bul("cari kodu", "hesap kodu", "kod"), "unvan": bul("unvan", "cari unvan", "cari adı", "ad"), "vkn": bul("vergi no", "vkn", "vkn/tckn", "tckn"),
         "iban": bul("iban", "ıban"), "tur": bul("tür", "cari türü", "tip")}
    if k["unvan"] is None:
        raise SystemExit(f"Cari listesinde Unvan sütunu gerekli. Başlıklar: {s[0]}")
    al = lambda r, a: r[k[a]] if k[a] is not None and k[a] < len(r) else None  # noqa: E731
    return [{"kod": str(al(r, "kod") or "").strip(), "unvan": str(al(r, "unvan")).strip(), "anahtar": unvan_anahtari(str(al(r, "unvan"))),
             "vkn": re.sub(r"\D", "", str(al(r, "vkn") or "")), "iban": re.sub(r"\s", "", str(al(r, "iban") or "")).upper(),
             "tur": kucuk(al(r, "tur"))} for r in s[1:] if al(r, "unvan")]


def kurallari_oku(yol: Path | None) -> list[dict]:
    if not yol:
        return []
    s = tablo_oku(yol)
    b = [kucuk(x) for x in s[0]]
    bul = lambda *a: next((i for i, x in enumerate(b) if x in a), None)  # noqa: E731
    k = {"ifade": bul("anahtar kelime", "ifade", "açıklama içerir"), "yon": bul("yön", "yon"), "hesap": bul("hesap kodu", "hesap"),
         "aciklama": bul("fiş açıklaması", "açıklama"), "min": bul("en az tutar", "min tutar"), "max": bul("en çok tutar", "max tutar")}
    if k["ifade"] is None or k["hesap"] is None:
        raise SystemExit(f"Kurallarda Anahtar Kelime ve Hesap Kodu gerekli. Başlıklar: {s[0]}")
    al = lambda r, a: r[k[a]] if k[a] is not None and k[a] < len(r) else None  # noqa: E731
    kurallar = []
    for i, r in enumerate(s[1:], 2):
        ifade = str(al(r, "ifade") or "").strip()
        if not ifade:
            continue
        kelimeler = [kucuk(x) for x in re.split(r"\s*\|\s*", ifade) if x.strip()]
        kurallar.append({"sira": i, "ifade": ifade, "desen": re.compile(r"(?<!\w)(?:" + "|".join(re.escape(x) for x in kelimeler) + r")"),
                         "yon": kucuk(al(r, "yon")) or "her ikisi", "hesap": str(al(r, "hesap")).strip(), "aciklama": str(al(r, "aciklama") or ifade),
                         "min": ek.sayi_coz(al(r, "min")) if al(r, "min") not in (None, "") else None,
                         "max": ek.sayi_coz(al(r, "max")) if al(r, "max") not in (None, "") else None})
    return kurallar


# ----------------------------------------------------------------------------
# Eşleştirme
# ----------------------------------------------------------------------------

IBAN = re.compile(r"\bTR\d{2}(?:\s?\d{4}){5}\s?\d{2}\b", re.I)
VKN = re.compile(r"(?<!\d)(\d{10,11})(?!\d)")


def cari_bul(aciklama: str, cariler: list[dict]) -> tuple[dict, str] | None:
    for m in IBAN.finditer(aciklama):
        iban = re.sub(r"\s", "", m.group(0)).upper()
        c = next((c for c in cariler if c["iban"] and c["iban"] == iban), None)
        if c:
            return c, "IBAN"
    for m in VKN.finditer(aciklama):
        c = next((c for c in cariler if c["vkn"] and c["vkn"] == m.group(1)), None)
        if c:
            return c, "VKN/TCKN"
    kelimeler = set(unvan_anahtari(aciklama))
    adaylar = []
    for c in cariler:
        if not c["anahtar"]:
            continue
        ortak = [w for w in c["anahtar"] if w in kelimeler]
        oran = len(ortak) / len(c["anahtar"])
        if ortak and (oran == 1 or (len(ortak) >= 2 and oran >= 0.6)):
            adaylar.append((oran, len(ortak), c))
    if adaylar:
        adaylar.sort(key=lambda x: (-x[0], -x[1]))
        if len(adaylar) == 1 or adaylar[0][:2] != adaylar[1][:2]:
            return adaylar[0][2], "Unvan"
    return None


def kural_bul(aciklama: str, tutar: float, kurallar: list[dict]) -> dict | None:
    a = kucuk(aciklama)
    yon = "giriş" if tutar > 0 else "çıkış"
    for k in kurallar:
        if k["yon"] not in ("her ikisi", "", yon) and not (k["yon"].startswith("gir") and yon == "giriş") and not (k["yon"].startswith("çık") and yon == "çıkış"):
            continue
        if k["min"] is not None and abs(tutar) < k["min"]:
            continue
        if k["max"] is not None and abs(tutar) > k["max"]:
            continue
        if k["desen"].search(a):
            return k
    return None


def oneriler_uret(hareketler, cariler, kurallar, banka_hesap: str, musteri_hesap: str, tedarikci_hesap: str) -> list[dict]:
    sonuc = []
    for h in hareketler:
        tutar = float(Decimal(str(h.tutar)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))
        if tutar == 0:
            continue
        yon = "Giriş" if tutar > 0 else "Çıkış"
        bulunan = cari_bul(h.aciklama, cariler)
        if bulunan:
            c, yol = bulunan
            varsayilan = musteri_hesap if tutar > 0 else tedarikci_hesap
            if c["tur"].startswith(("müşteri", "alıcı")) and tutar < 0:
                yol += " · müşteriye ödeme (iade?)"
            if c["tur"].startswith(("tedarikçi", "satıcı")) and tutar > 0:
                yol += " · tedarikçiden tahsilat (iade?)"
            karsi = c["kod"] or varsayilan
            aciklama = f"{c['unvan']} {'tahsilat' if tutar > 0 else 'ödeme'}"
            kaynak, guven = f"Cari ({yol})", "Yüksek" if yol.startswith(("IBAN", "VKN")) else "Orta"
        else:
            k = kural_bul(h.aciklama, tutar, kurallar)
            if k:
                karsi, aciklama, kaynak, guven = k["hesap"], k["aciklama"], f"Kural {k['sira']}: '{k['ifade']}'", "Orta"
            else:
                karsi, aciklama, kaynak, guven = "", "", "—", "Manuel"
        sonuc.append({"h": h, "tutar": tutar, "yon": yon, "karsi": karsi, "aciklama": aciklama or h.aciklama, "kaynak": kaynak, "guven": guven,
                      "borc_hesap": banka_hesap if tutar > 0 else karsi, "alacak_hesap": karsi if tutar > 0 else banka_hesap})
    return sonuc


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
GUVEN_DOLGU = {"Yüksek": PatternFill("solid", fgColor="E3F5E1"), "Orta": PatternFill("solid", fgColor="FFF4CE"),
               "Manuel": PatternFill("solid", fgColor="FDE2E1")}
PARA = "#,##0.00"
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, b, g):
    ws.append(b)
    for c in ws[ws.max_row]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI
    for j, w in enumerate(g, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def calistir(ekstre: Path, cikti: Path, cari_yolu: Path | None = None, kural_yolu: Path | None = None, banka_hesap: str = "102.01",
             musteri_hesap: str = "120", tedarikci_hesap: str = "320", fis_baslangic: int = 1) -> dict:
    hareketler = ek.dosya_oku(ekstre, "banka")
    cariler, kurallar = carileri_oku(cari_yolu), kurallari_oku(kural_yolu)
    oneriler = oneriler_uret(hareketler, cariler, kurallar, banka_hesap, musteri_hesap, tedarikci_hesap)

    wb = Workbook()
    ws = wb.active
    ws.title = "Hareket Eşleşmeleri"
    _baslik(ws, ["Satır", "Tarih", "Ekstre Açıklaması", "Yön", "Tutar", "Karşı Hesap", "Fiş Açıklaması", "Eşleşme", "Güven", "Onay"],
            (7, 12, 50, 8, 14, 14, 40, 34, 9, 9))
    for o in oneriler:
        ws.append([o["h"].sira, o["h"].tarih, o["h"].aciklama, o["yon"], abs(o["tutar"]), o["karsi"], o["aciklama"], o["kaynak"], o["guven"], ""])
        ws.cell(ws.max_row, 2).number_format = "DD.MM.YYYY"
        ws.cell(ws.max_row, 5).number_format = PARA
        ws.cell(ws.max_row, 9).fill = GUVEN_DOLGU[o["guven"]]
        ws.cell(ws.max_row, 10).fill = GUVEN_DOLGU["Orta"]
    ws.auto_filter.ref = ws.dimensions

    f = wb.create_sheet("Fiş Aktarım")
    _baslik(f, ["Fiş No", "Tarih", "Hesap Kodu", "Açıklama", "Borç", "Alacak", "Ekstre Satırı"], (8, 12, 16, 50, 14, 14, 12))
    fis = fis_baslangic
    for o in oneriler:
        if o["guven"] == "Manuel":
            continue
        for hesap, borc, alacak in ((o["borc_hesap"], abs(o["tutar"]), None), (o["alacak_hesap"], None, abs(o["tutar"]))):
            f.append([fis, o["h"].tarih, hesap, o["aciklama"], borc, alacak, o["h"].sira])
            f.cell(f.max_row, 2).number_format = "DD.MM.YYYY"
            f.cell(f.max_row, 5).number_format = f.cell(f.max_row, 6).number_format = PARA
        fis += 1

    m = wb.create_sheet("Manuel")
    _baslik(m, ["Satır", "Tarih", "Ekstre Açıklaması", "Yön", "Tutar", "Önerilen Karşı Hesap (doldurun)"], (7, 12, 60, 8, 14, 30))
    for o in oneriler:
        if o["guven"] == "Manuel":
            m.append([o["h"].sira, o["h"].tarih, o["h"].aciklama, o["yon"], abs(o["tutar"]), ""])
            m.cell(m.max_row, 2).number_format = "DD.MM.YYYY"
            m.cell(m.max_row, 5).number_format = PARA
            m.cell(m.max_row, 6).fill = GUVEN_DOLGU["Orta"]

    say = Counter(o["guven"] for o in oneriler)
    b = wb.create_sheet("Bilgi")
    for s in [["Sonuç", f"{len(oneriler)} hareket · yüksek güven {say['Yüksek']} · orta {say['Orta']} · manuel {say['Manuel']}"],
              ["Banka hesabı", banka_hesap],
              ["Cari eşleşme", f"Sıra: IBAN → VKN/TCKN → unvan (şirket türü ekleri atılır, ayırt edici kelimelerin tamamı veya en az 2'si ve %60'ı). "
                               f"Giriş: müşteri ({musteri_hesap}) alacak; çıkış: tedarikçi ({tedarikci_hesap}) borç; cari kodu varsa o kullanılır"],
              ["Kurallar", "Kurallar dosyasında yukarıdan aşağı ilk uyan kural uygulanır; 'a | b' biçimi 'a veya b' demektir. Örnek kurallar "
                           "genel uygulamayı yansıtır; hesap kodlarını kendi hesap planınıza ve mali müşavirinizin tercihine göre düzenleyin"],
              ["Fiş", "Giriş: banka BORÇ / karşı hesap ALACAK · Çıkış: karşı hesap BORÇ / banka ALACAK. KDV, stopaj ve BSMV ayrıştırması yapılmaz"],
              ["Onay", "Öneriler taslaktır; muhasebeleştirmeden önce kontrol edin"],
              ["Oluşturulma", datetime.now().strftime("%d.%m.%Y %H:%M")]]:
        b.append(s)
    b.column_dimensions["A"].width = 16
    b.column_dimensions["B"].width = 130
    for c in b["B"]:
        c.alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    return {"oneriler": oneriler, "sayac": say}


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description="Banka ekstresi hareketlerinden muhasebe fişi önerisi üretir.")
    ap.add_argument("--ekstre", type=Path, default=BURASI / "ornek_veri" / "banka_ekstresi.csv",
                    help="Banka ekstresi (.xlsx/.csv): Tarih, Açıklama ve Tutar (veya Borç + Alacak)")
    ap.add_argument("--cariler", type=Path, help="Cari kartlar (.xlsx/.csv): Unvan; Cari Kodu, Vergi No, IBAN, Tür (Müşteri/Tedarikçi)")
    ap.add_argument("--kurallar", type=Path, help="Hesap kuralları (.csv/.xlsx): Anahtar Kelime, Yön, Hesap Kodu, Fiş Açıklaması, En Az/En Çok Tutar")
    ap.add_argument("--banka-hesap", default="102.01", help="Banka hesap kodu (varsayılan 102.01)")
    ap.add_argument("--musteri-hesap", default="120", help="Cari kodu olmayan müşteri tahsilatı hesabı (varsayılan 120)")
    ap.add_argument("--tedarikci-hesap", default="320", help="Cari kodu olmayan tedarikçi ödemesi hesabı (varsayılan 320)")
    ap.add_argument("--fis-no", type=int, default=1, help="İlk fiş numarası")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "muhasebe_fisi_onerisi.xlsx")
    a = ap.parse_args(argv)
    if a.ekstre == BURASI / "ornek_veri" / "banka_ekstresi.csv":
        a.cariler = a.cariler or BURASI / "cari_kartlar_ornek.csv"
        a.kurallar = a.kurallar or BURASI / "kurallar_ornek.csv"
    s = calistir(a.ekstre, a.cikti, a.cariler, a.kurallar, a.banka_hesap, a.musteri_hesap, a.tedarikci_hesap, a.fis_no)
    say = s["sayac"]
    print(f"[OK] {len(s['oneriler'])} hareket · yüksek güven {say['Yüksek']} · orta {say['Orta']} · manuel {say['Manuel']}")
    for o in s["oneriler"]:
        if o["guven"] == "Manuel":
            print(f"[?] {o['h'].tarih:%d.%m.%Y} {o['h'].aciklama[:50]:<50} {o['tutar']:>12,.2f}".replace(",", "X").replace(".", ",").replace("X", "."))
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
