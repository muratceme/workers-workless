import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))

import main  # noqa: E402

ORNEK = KLASOR / "ornek_veri" / "banka_ekstresi.csv"
CARI = KLASOR / "cari_kartlar_ornek.csv"
KURAL = KLASOR / "kurallar_ornek.csv"


class EslesmeTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.cariler = main.carileri_oku(CARI)
        cls.kurallar = main.kurallari_oku(KURAL)

    def test_iban_ve_vkn(self):
        c, yol = main.cari_bul("Giden EFT TR33 0006 1005 1978 6457 8413 26 fatura ödemesi", self.cariler)
        self.assertEqual((c["kod"], yol), ("320.01.004", "IBAN"))
        c, yol = main.cari_bul("Gelen havale 6620045671", self.cariler)
        self.assertEqual((c["kod"], yol), ("120.01.010", "VKN/TCKN"))

    def test_unvan(self):
        self.assertEqual(main.cari_bul("Gelen EFT ÖRNEK PERAKENDE MAĞAZA", self.cariler)[0]["kod"], "120.01.010")
        self.assertIsNone(main.cari_bul("Gelen EFT Deneme", [{"kod": "1", "unvan": "Deneme Tekstil", "anahtar": ["deneme"], "vkn": "", "iban": "", "tur": ""},
                                                            {"kod": "2", "unvan": "Deneme Gıda", "anahtar": ["deneme"], "vkn": "", "iban": "", "tur": ""}]))
        self.assertIsNone(main.cari_bul("Gelen EFT Mehmet", self.cariler))

    def test_kural_yon_ve_tutar(self):
        self.assertEqual(main.kural_bul("Mevduat faizi", 100, self.kurallar)["hesap"], "642")
        self.assertIsNone(main.kural_bul("Kredi faizi tahsil", -100, self.kurallar))          # faiz kuralı yalnız giriş
        self.assertIsNone(main.kural_bul("Danışmanlık ücreti", -20000, self.kurallar))        # masraf kuralı en çok 5.000
        self.assertEqual(main.kural_bul("Muhtasar ve KDV ödemesi", -5000, self.kurallar)["hesap"], "360")


class UctanUcaTest(unittest.TestCase):
    def test_ornek_ve_fis_dengesi(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "f.xlsx"
            s = main.calistir(ORNEK, cikti, CARI, KURAL, "102.01.001")
            o = {x["h"].sira: x for x in s["oneriler"]}
            self.assertEqual((o[1]["borc_hesap"], o[1]["alacak_hesap"]), ("102.01.001", "120.01.001"))   # tahsilat
            self.assertEqual((o[2]["borc_hesap"], o[2]["alacak_hesap"]), ("320.01.004", "102.01.001"))   # ödeme
            self.assertEqual(s["sayac"]["Manuel"], 0)
            satirlar = list(load_workbook(cikti)["Fiş Aktarım"].iter_rows(min_row=2, values_only=True))
            self.assertEqual(len(satirlar), 30)
            self.assertAlmostEqual(sum(r[4] or 0 for r in satirlar), sum(r[5] or 0 for r in satirlar))

    def test_eslesmeyen_manuel(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t / "e.csv").write_text("Tarih;Açıklama;Tutar\n01.10.2026;Bilinmeyen gönderen;500\n", encoding="utf-8")
            s = main.calistir(t / "e.csv", t / "f.xlsx", CARI, KURAL)
        self.assertEqual(s["oneriler"][0]["guven"], "Manuel")


if __name__ == "__main__":
    unittest.main()
