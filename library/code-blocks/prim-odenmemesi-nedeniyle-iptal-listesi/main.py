"""
Prim Ödenmemesi Nedeniyle İptal Listesi — Gerisi Bende kod bloğu
Tahsilat Yönetimi › Tahsilat Uzmanı

Vadesi geçtiği hâlde ödenmeyen prim taksitlerini süre kurallarına göre sınıflar; iptal adayı poliçeleri listeler.
  Gecikmiş taksit              = kalan tutar (tutar − ödenen) > 0 ve vade < durum tarihi
  İlk taksit (peşinat)         : cayma süresi sonu = vade + --cayma-ay (3) ay
      Peşinat ödenmedi             = süre içinde;   Peşinat ödenmedi (süre doldu) = durum tarihi > süre sonu
  İzleyen taksitler            : ihtar süresi sonu = ihtarın tebliğ tarihi + --ihtar-sure (10) gün
      İhtar gönderilecek           = ihtar tebliğ tarihi boş
      İhtar süresi işliyor         = durum tarihi ≤ süre sonu
      İptal adayı                  = durum tarihi > süre sonu ve taksit hâlâ ödenmemiş
  Poliçenin durumu             = taksitlerinin en ağır durumu (İptal adayı > Peşinat süre doldu > Peşinat ödenmedi > İhtar süresi işliyor > İhtar gönderilecek)
Süreler çalışma varsayımıdır; branşın genel ve özel şartlarına, zorunlu sigortalardaki özel kurallara göre değişir. Kod iptal kararı vermez.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
import calendar

SUTUN = {"police": ("police no", "police numarasi", "police"), "sigortali": ("sigortali", "sigorta ettiren", "musteri"), "brans": ("brans", "urun"), "no": ("taksit no", "taksit", "taksit sirasi"),
         "vade": ("vade", "vade tarihi", "taksit vadesi"), "tutar": ("tutar", "taksit tutari", "prim"), "odenen": ("odenen", "odenen tutar", "tahsil edilen"), "ihtar": ("ihtar teblig tarihi", "ihtar tarihi", "teblig tarihi")}
DURUM = ("İptal adayı", "Peşinat ödenmedi (süre doldu)", "Peşinat ödenmedi", "İhtar süresi işliyor", "İhtar gönderilecek")
DURUM_RENK = {"İptal adayı": KIRMIZI, "Peşinat ödenmedi (süre doldu)": KIRMIZI, "Peşinat ödenmedi": SARI, "İhtar süresi işliyor": MAVI, "İhtar gönderilecek": SARI, "Vadesi gelmedi": None, "Ödendi": YESIL}
ONEM = {"İptal adayı": "Yüksek", "Peşinat ödenmedi (süre doldu)": "Yüksek", "Peşinat ödenmedi": "Orta", "İhtar süresi işliyor": "Bilgi", "İhtar gönderilecek": "Orta"}


def ay_ekle(t: date, n: int) -> date:
    y, m = divmod(t.year * 12 + t.month - 1 + n, 12)
    return date(y, m + 1, min(t.day, calendar.monthrange(y, m + 1)[1]))


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("police", "vade", "tutar")):
        police = metin(r["police"])
        if not police:
            continue
        vade, tutar, odenen, no = tarih(r["vade"]), para(r["tutar"]), para(r["odenen"]) if metin(r["odenen"]) else Decimal(0), para(r["no"]) if metin(r["no"]) else None
        if vade is None or tutar is None or tutar <= 0 or odenen is None or odenen < 0 or (metin(r["no"]) and (no is None or no < 1 or no != no.to_integral_value())):
            uy.append(uyari("Yüksek", "Taksit okunamadı", police, f"Satır {r['_satir']}: vade '{metin(r['vade'])}', tutar '{metin(r['tutar'])}', ödenen '{metin(r['odenen'])}' ya da taksit no '{metin(r['no'])}' okunamadı. Satır alınmadı."))
            continue
        if metin(r["ihtar"]) and tarih(r["ihtar"]) is None:
            uy.append(uyari("Orta", "Tarih okunamadı", police, f"İhtar tebliğ tarihi '{metin(r['ihtar'])}' okunamadı; ihtar gönderilmemiş sayıldı."))
        l.append({"police": police, "sigortali": metin(r["sigortali"]), "brans": metin(r["brans"]) or "(branş yok)", "no": int(no) if no is not None else None, "vade": vade, "tutar": tutar, "odenen": odenen, "ihtar": tarih(r["ihtar"])})
    return l, uy


def degerlendir(taksitler: list[dict], bugun: date, ihtar_sure: int = 10, cayma_ay: int = 3, gun: int = 15) -> dict:
    if not taksitler:
        raise ValueError("Değerlendirilecek taksit yok.")
    uy, grup = [], {}
    for t in taksitler:
        grup.setdefault(katla(t["police"]), []).append(t)
    policeler = []
    for g in grup.values():
        g.sort(key=lambda t: (t["vade"], t["no"] or 0))
        numarali = all(t["no"] is not None for t in g)
        for i, t in enumerate(g, 1):
            t["sira"] = t["no"] if numarali else i
            t["kalan"] = max(t["tutar"] - t["odenen"], Decimal(0))
            t["gecikme"] = t["son"] = t["sure_kalan"] = None
            if t["odenen"] > t["tutar"]:
                uy.append(uyari("Orta", "Ödenen tutar taksitten büyük", f"{t['police']} · {t['sira']}. taksit", f"Taksit {t['tutar']:,.2f}, ödenen {t['odenen']:,.2f}. Fazla ödeme başka taksite mahsup edilmedi; kayıtları kontrol edin."))
            if t["kalan"] == 0:
                t["durum"] = "Ödendi"
            elif t["vade"] >= bugun:
                t["durum"] = "Vadesi gelmedi"
            else:
                t["gecikme"] = (bugun - t["vade"]).days
                if t["sira"] == 1:
                    t["son"] = ay_ekle(t["vade"], cayma_ay)
                    t["durum"] = "Peşinat ödenmedi (süre doldu)" if bugun > t["son"] else "Peşinat ödenmedi"
                elif t["ihtar"] is None:
                    t["durum"] = "İhtar gönderilecek"
                else:
                    if t["ihtar"] < t["vade"]:
                        uy.append(uyari("Orta", "İhtar tarihi vadeden önce", f"{t['police']} · {t['sira']}. taksit", f"Vade {t['vade']:%d.%m.%Y}, ihtar tebliği {t['ihtar']:%d.%m.%Y}. Tarihleri kontrol edin."))
                    t["son"] = t["ihtar"] + timedelta(days=ihtar_sure)
                    t["durum"] = "İptal adayı" if bugun > t["son"] else "İhtar süresi işliyor"
                if t["son"]:
                    t["sure_kalan"] = (t["son"] - bugun).days
        gec = [t for t in g if t["gecikme"] is not None]
        if not gec:
            continue
        agir = min(gec, key=lambda t: (DURUM.index(t["durum"]), t["vade"]))
        p = {"police": g[0]["police"], "sigortali": g[0]["sigortali"], "brans": g[0]["brans"], "durum": agir["durum"], "taksit": len(gec), "tutar": sum((t["kalan"] for t in gec), Decimal(0)), "en_eski": gec[0]["vade"], "gecikme": gec[0]["gecikme"],
             "son": agir["son"], "sure_kalan": agir["sure_kalan"], "ayrinti": "; ".join(f"{t['sira']}. taksit ({t['vade']:%d.%m.%Y}, kalan {t['kalan']:,.2f}): {t['durum']}" for t in gec),
             "gelecek": sum((t["kalan"] for t in g if t["durum"] == "Vadesi gelmedi"), Decimal(0))}
        policeler.append(p)
        kim = f"{p['police']} · {p['sigortali']}" if p["sigortali"] else p["police"]
        if p["durum"] == "İptal adayı":
            uy.append(uyari("Yüksek", "İptal adayı", kim, f"İhtar süresi {p['son']:%d.%m.%Y} tarihinde doldu, taksit ödenmedi. Gecikmiş tutar {p['tutar']:,.2f}. {p['ayrinti']}. İptal işleminden önce ihtarın usulüne uygun tebliğ edildiğini doğrulayın."))
        elif p["durum"] == "Peşinat ödenmedi (süre doldu)":
            uy.append(uyari("Yüksek", "Peşinat ödenmedi, cayma süresi doldu", kim, f"İlk taksitin vadesi {p['en_eski']:%d.%m.%Y}; {cayma_ay} aylık süre {p['son']:%d.%m.%Y} tarihinde doldu. Poliçenin hukuki durumunu değerlendirin. {p['ayrinti']}."))
        elif p["durum"] == "Peşinat ödenmedi":
            yakin = p["sure_kalan"] <= gun
            uy.append(uyari("Yüksek" if yakin else "Orta", "Cayma süresi dolmak üzere" if yakin else "Peşinat ödenmedi", kim, f"İlk taksit {p['gecikme']} gündür ödenmedi; cayma süresi {p['son']:%d.%m.%Y} tarihinde doluyor ({p['sure_kalan']} gün). "
                            f"Gecikmiş tutar {p['tutar']:,.2f}."))
        elif p["durum"] == "İhtar gönderilecek":
            uy.append(uyari("Orta", "İhtar gönderilecek", kim, f"{p['taksit']} taksit gecikmiş ({p['tutar']:,.2f}); en eski vade {p['en_eski']:%d.%m.%Y} ({p['gecikme']} gün). {p['ayrinti']}."))
        else:
            uy.append(uyari("Bilgi", "İhtar süresi işliyor", kim, f"Süre {p['son']:%d.%m.%Y} tarihinde doluyor ({p['sure_kalan']} gün). Gecikmiş tutar {p['tutar']:,.2f}."))
    policeler.sort(key=lambda p: (DURUM.index(p["durum"]), -p["gecikme"], katla(p["police"])))
    durumlar = [{"durum": d, "police": sum(1 for p in policeler if p["durum"] == d), "tutar": sum((p["tutar"] for p in policeler if p["durum"] == d), Decimal(0))} for d in DURUM]
    brans = {}
    for p in policeler:
        b = brans.setdefault(katla(p["brans"]), {"ad": p["brans"], "police": 0, "tutar": Decimal(0), "aday": 0})
        b["police"], b["tutar"], b["aday"] = b["police"] + 1, b["tutar"] + p["tutar"], b["aday"] + (p["durum"] == "İptal adayı")
    return {"taksitler": taksitler, "policeler": policeler, "durumlar": durumlar, "branslar": sorted(brans.values(), key=lambda b: (-b["tutar"], katla(b["ad"]))),
            "ozet": {"police": len(grup), "gecikmis": len(policeler), "aday": durumlar[0]["police"], "tutar": sum((p["tutar"] for p in policeler), Decimal(0))}, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, bugun: date, ihtar_sure: int, cayma_ay: int) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Poliçeler"
    _baslik(ws, ["Durum", "Poliçe No", "Sigortalı", "Branş", "Gecikmiş Taksit", "Gecikmiş Tutar", "En Eski Vade", "Gecikme (gün)", "Süre Sonu", "Süreye Kalan (gün)", "Vadesi Gelmemiş Tutar", "Taksit Ayrıntısı", "Yapılan İşlem (siz doldurun)"],
            (28, 14, 28, 14, 9, 15, 12, 10, 12, 10, 15, 70, 34))
    for p in r["policeler"]:
        _satir(ws, [p["durum"], p["police"], p["sigortali"], p["brans"], p["taksit"], p["tutar"], p["en_eski"], p["gecikme"], p["son"], p["sure_kalan"], p["gelecek"] or None, p["ayrinti"], None], para_s=(6, 11), tarih_s=(7, 9),
               dolgu={1: DURUM_RENK[p["durum"]]}, sar=(12,))
    ws = wb.create_sheet("Özet")
    o = r["ozet"]
    _baslik(ws, ["Gösterge", "Değer"], (40, 18))
    for etiket, deger, para_mi in (("Dosyadaki poliçe", o["police"], False), ("Gecikmiş taksiti olan poliçe", o["gecikmis"], False), ("İptal adayı poliçe", o["aday"], False), ("Gecikmiş prim tutarı", o["tutar"], True)):
        _satir(ws, [etiket, deger], para_s=(2,) if para_mi else ())
    ws.append([])
    _baslik(ws, ["Durum", "Poliçe", "Gecikmiş Tutar"], dondur="A2")
    for d in r["durumlar"]:
        _satir(ws, [d["durum"], d["police"], d["tutar"]], para_s=(3,), dolgu={1: DURUM_RENK[d["durum"]] if d["police"] else None})
    ws.column_dimensions["C"].width = 18
    ws.append([])
    _baslik(ws, ["Branş", "Poliçe", "Gecikmiş Tutar", "İptal Adayı"], dondur="A2")
    for b in r["branslar"]:
        _satir(ws, [b["ad"], b["police"], b["tutar"], b["aday"] or None], para_s=(3,))
    ws.append([])
    ws.append([f"Durum tarihi: {bugun:%d.%m.%Y} · ihtar süresi {ihtar_sure} gün · peşinat için cayma süresi {cayma_ay} ay (çalışma varsayımı)"])
    ws = wb.create_sheet("Taksitler")
    _baslik(ws, ["Poliçe No", "Sigortalı", "Branş", "Taksit", "Vade", "Tutar", "Ödenen", "Kalan", "Durum", "Gecikme (gün)", "İhtar Tebliğ", "Süre Sonu"], (14, 28, 14, 7, 12, 14, 14, 14, 28, 10, 12, 12))
    for t in sorted(r["taksitler"], key=lambda t: (katla(t["police"]), t["vade"], t["sira"])):
        _satir(ws, [t["police"], t["sigortali"], t["brans"], t["sira"], t["vade"], t["tutar"], t["odenen"], t["kalan"], t["durum"], t["gecikme"], t["ihtar"], t["son"]], para_s=(6, 7, 8), tarih_s=(5, 11, 12), dolgu={9: DURUM_RENK[t["durum"]]})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, ihtar_sure: int = 10, cayma_ay: int = 3, gun: int = 15) -> dict:
    taksitler, uy = oku(yol)
    r = degerlendir(taksitler, bugun, ihtar_sure, cayma_ay, gun)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, bugun, ihtar_sure, cayma_ay)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Taksiti ödenmeyen poliçeleri süre kurallarına göre sınıflar; iptal adaylarını listeler.")
    p.add_argument("--taksitler", type=Path, help="Poliçe No, Sigortalı, Branş, Taksit No, Vade, Tutar, Ödenen, İhtar Tebliğ Tarihi")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--ihtar-sure", type=int, default=10, help="İzleyen taksitlerde ihtarın tebliğinden sonra tanınan süre, gün (varsayılan 10)")
    p.add_argument("--cayma-ay", type=int, default=3, help="İlk taksit ödenmediğinde vadeden başlayan cayma süresi, ay (varsayılan 3)")
    p.add_argument("--gun", type=int, default=15, help="Cayma süresinin dolmasına kaç gün kala uyarı yükseltilsin (varsayılan 15)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "iptal_adayi_policeler.xlsx")
    a = p.parse_args(argv)
    ornek = a.taksitler is None
    yol = ORNEK / "taksitler.csv" if ornek else a.taksitler
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    if bugun is None or a.ihtar_sure < 1 or a.cayma_ay < 1 or a.gun < 0:
        print("[X] --tarih GG.AA.YYYY; --ihtar-sure ve --cayma-ay en az 1; --gun eksi olamaz")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, bugun, a.ihtar_sure, a.cayma_ay, a.gun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {o['police']} poliçe · gecikmiş taksiti olan {o['gecikmis']} · iptal adayı {o['aday']} · gecikmiş prim {o['tutar']:,.2f}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
