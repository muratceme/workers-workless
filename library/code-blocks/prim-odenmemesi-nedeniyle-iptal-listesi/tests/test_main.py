import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)


def t(police, no, vade, tutar=1000, odenen=0, ihtar=None):
    return {"police": police, "sigortali": "", "brans": "Kasko", "no": no, "vade": vade, "tutar": D(tutar), "odenen": D(odenen), "ihtar": ihtar}


class DegerlendirTesti(unittest.TestCase):
    def test_durumlar(self):
        l = [t("aday", 2, date(2026, 9, 1), ihtar=date(2026, 10, 1)), t("işliyor", 2, date(2026, 9, 1), ihtar=date(2026, 10, 2)), t("ihtarsız", 3, date(2026, 10, 11)), t("bugün", 2, date(2026, 10, 12)),
             t("peşin", 1, date(2026, 7, 12)), t("peşin_doldu", 1, date(2026, 7, 11)), t("ödendi", 2, date(2026, 9, 1), odenen=1000), t("kısmi", 2, date(2026, 9, 1), odenen=400), t("fazla", 2, date(2026, 9, 1), odenen=1200),
             t("erken_ihtar", 2, date(2026, 9, 10), ihtar=date(2026, 9, 5))]
        r = main.degerlendir(l, BUGUN)
        self.assertEqual({x["police"]: (x["durum"], x["kalan"], x["gecikme"], x["son"], x["sure_kalan"]) for x in r["taksitler"]},
                         {"aday": ("İptal adayı", D(1000), 41, date(2026, 10, 11), -1), "işliyor": ("İhtar süresi işliyor", D(1000), 41, date(2026, 10, 12), 0), "ihtarsız": ("İhtar gönderilecek", D(1000), 1, None, None),
                          "bugün": ("Vadesi gelmedi", D(1000), None, None, None), "peşin": ("Peşinat ödenmedi", D(1000), 92, date(2026, 10, 12), 0), "peşin_doldu": ("Peşinat ödenmedi (süre doldu)", D(1000), 93, date(2026, 10, 11), -1),
                          "ödendi": ("Ödendi", D(0), None, None, None), "kısmi": ("İhtar gönderilecek", D(600), 41, None, None), "fazla": ("Ödendi", D(0), None, None, None),
                          "erken_ihtar": ("İptal adayı", D(1000), 32, date(2026, 9, 15), -27)})
        self.assertEqual([(p["police"], p["durum"]) for p in r["policeler"]],
                         [("aday", "İptal adayı"), ("erken_ihtar", "İptal adayı"), ("peşin_doldu", "Peşinat ödenmedi (süre doldu)"), ("peşin", "Peşinat ödenmedi"), ("işliyor", "İhtar süresi işliyor"), ("kısmi", "İhtar gönderilecek"),
                          ("ihtarsız", "İhtar gönderilecek")])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         sorted([("Yüksek", "İptal adayı", "aday"), ("Yüksek", "İptal adayı", "erken_ihtar"), ("Orta", "İhtar tarihi vadeden önce", "erken_ihtar · 2. taksit"), ("Yüksek", "Peşinat ödenmedi, cayma süresi doldu", "peşin_doldu"),
                                 ("Yüksek", "Cayma süresi dolmak üzere", "peşin"), ("Bilgi", "İhtar süresi işliyor", "işliyor"), ("Orta", "İhtar gönderilecek", "kısmi"), ("Orta", "İhtar gönderilecek", "ihtarsız"),
                                 ("Orta", "Ödenen tutar taksitten büyük", "fazla · 2. taksit")]))
        self.assertEqual((r["ozet"], [(d["durum"], d["police"], d["tutar"]) for d in r["durumlar"]]),
                         ({"police": 10, "gecikmis": 7, "aday": 2, "tutar": D(6600)},
                          [("İptal adayı", 2, D(2000)), ("Peşinat ödenmedi (süre doldu)", 1, D(1000)), ("Peşinat ödenmedi", 1, D(1000)), ("İhtar süresi işliyor", 1, D(1000)), ("İhtar gönderilecek", 2, D(1600))]))
        d = lambda l, **kw: [x["durum"] for x in main.degerlendir(l, BUGUN, **kw)["taksitler"]]   # noqa: E731
        self.assertEqual((d([t("a", 2, date(2026, 9, 1), ihtar=date(2026, 10, 1))], ihtar_sure=11), d([t("a", 1, date(2026, 7, 11))], cayma_ay=4)), (["İhtar süresi işliyor"], ["Peşinat ödenmedi"]))
        self.assertEqual((main.ay_ekle(date(2026, 11, 30), 3), main.ay_ekle(date(2026, 1, 31), 1), main.ay_ekle(date(2026, 10, 15), 3)), (date(2027, 2, 28), date(2026, 2, 28), date(2027, 1, 15)))
        u = lambda **kw: [x["tur"] for x in main.degerlendir([t("a", 1, date(2026, 8, 1))], BUGUN, **kw)["uyarilar"]]   # noqa: E731
        self.assertEqual((u(), u(gun=20)), (["Peşinat ödenmedi"], ["Cayma süresi dolmak üzere"]))
        with self.assertRaises(ValueError):
            main.degerlendir([], BUGUN)

    def test_police_ozeti(self):
        l = [t("P", None, date(2026, 9, 10)), t("p", None, date(2026, 6, 10), odenen=1000), t("P", None, date(2026, 8, 10), 500, 100, date(2026, 9, 1)), t("P", None, date(2026, 12, 10), 700),
             t("Q", 5, date(2026, 6, 1), odenen=1000), t("Q", 6, date(2026, 11, 1))]
        r = main.degerlendir(l, BUGUN)
        self.assertEqual([(x["vade"], x["sira"], x["durum"]) for x in sorted((x for x in r["taksitler"] if x["police"].upper() == "P"), key=lambda x: x["vade"])],
                         [(date(2026, 6, 10), 1, "Ödendi"), (date(2026, 8, 10), 2, "İptal adayı"), (date(2026, 9, 10), 3, "İhtar gönderilecek"), (date(2026, 12, 10), 4, "Vadesi gelmedi")])
        self.assertEqual(len(r["policeler"]), 1)
        p = r["policeler"][0]
        self.assertEqual((p["durum"], p["taksit"], p["tutar"], p["en_eski"], p["gecikme"], p["son"], p["sure_kalan"], p["gelecek"], p["ayrinti"]),
                         ("İptal adayı", 2, D(1400), date(2026, 8, 10), 63, date(2026, 9, 11), -31, D(700), "2. taksit (10.08.2026, kalan 400.00): İptal adayı; 3. taksit (10.09.2026, kalan 1,000.00): İhtar gönderilecek"))
        self.assertEqual((r["ozet"], [(b["ad"], b["police"], b["tutar"], b["aday"]) for b in r["branslar"]]), ({"police": 2, "gecikmis": 1, "aday": 1, "tutar": D(1400)}, [("Kasko", 1, D(1400), 1)]))


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "taksitler.csv", cikti, BUGUN)
            self.assertEqual([(p["police"], p["durum"], p["taksit"], p["tutar"], p["gecikme"], p["son"], p["sure_kalan"]) for p in r["policeler"]],
                             [("P-1001", "İptal adayı", 2, D(24000), 63, date(2026, 10, 5), -7), ("P-1004", "Peşinat ödenmedi (süre doldu)", 1, D(20000), 103, date(2026, 10, 1), -11),
                              ("P-1002", "Peşinat ödenmedi", 1, D(30000), 27, date(2026, 12, 15), 64), ("P-1003", "İhtar süresi işliyor", 1, D(5000), 41, date(2026, 10, 15), 3), ("P-1006", "İhtar gönderilecek", 1, D(9000), 22, None, None)])
            self.assertEqual(r["ozet"], {"police": 6, "gecikmis": 5, "aday": 1, "tutar": D(88000)})
            self.assertEqual([(b["ad"], b["police"], b["tutar"]) for b in r["branslar"]], [("Kasko", 2, D(33000)), ("Yangın", 1, D(30000)), ("İşyeri", 1, D(20000)), ("Sağlık", 1, D(5000))])
            self.assertEqual([(u["onem"], u["tur"], u["kim"][:6]) for u in r["uyarilar"]],
                             [("Yüksek", "Peşinat ödenmedi, cayma süresi doldu", "P-1004"), ("Yüksek", "İptal adayı", "P-1001"), ("Orta", "Peşinat ödenmedi", "P-1002"), ("Orta", "İhtar gönderilecek", "P-1006"), ("Bilgi", "İhtar süresi işliyor", "P-1003")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Poliçeler", "Özet", "Taksitler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "t.csv"
            y.write_text("Poliçe;Müşteri;Ürün;Taksit;Vade Tarihi;Taksit Tutarı;Tahsil Edilen;İhtar Tarihi\nP1;Ali;Kasko;2;01.09.2026;1.250,50;250,50;x\nP2;;;;02.09.2026;500;;05.10.2026\nP3;;;;x;500;;\nP4;;;;02.09.2026;0;;\n"
                         "P5;;;1,5;02.09.2026;10;;\nP6;;;;02.09.2026;10;-1;\n;;;;02.09.2026;10;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["police"], x["sigortali"], x["brans"], x["no"], x["vade"], x["tutar"], x["odenen"], x["ihtar"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("P1", "Ali", "Kasko", 2, date(2026, 9, 1), D("1250.50"), D("250.50"), None), ("P2", "", "(branş yok)", None, date(2026, 9, 2), D(500), D(0), date(2026, 10, 5))],
                              [("Tarih okunamadı", "P1"), ("Taksit okunamadı", "P3"), ("Taksit okunamadı", "P4"), ("Taksit okunamadı", "P5"), ("Taksit okunamadı", "P6")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--taksitler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "12.10.2026", "--ihtar-sure", "15", "--cayma-ay", "2", "--gun", "30"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--ihtar-sure", "0"], ["--cayma-ay", "0"], ["--gun", "-1"])], [2] * 4)
            self.assertEqual(main.main(["--taksitler", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Poliçe No;Vade\nA;01.09.2026\n", "Poliçe No;Vade;Tutar\n;01.09.2026;5\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
