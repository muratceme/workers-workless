"""
Rakip Fiyat Karşılaştırması — Gerisi Bende kod bloğu
Satın Alma ve Kategori Yönetimi › Kategori Uzmanı
Rakip fiyat taramalarını ürün eşleştirmesiyle karşılaştırıp fiyat endeksi çıkarır.
  Eşleştirme      = taramadaki ürün kodu bizim ürün koduna ya da barkoda eşitse
  Geçerli fiyat   = her ürün-rakip çiftinin en son (kampanyasız) tarama fiyatı; --eski (30) günden eski taramalar dışarıda
  Ürün endeksi    = bizim fiyat ÷ rakiplerin ortalama fiyatı × 100   (100 üstü: rakiplerden pahalıyız)
  Sepet endeksi   = Σ(bizim fiyat × ağırlık) ÷ Σ(rakip fiyatı × ağırlık) × 100 — rakip ve kategori bazında,
                    yalnız iki tarafta da fiyatı olan ürünlerle; ağırlık = satış adedi (yazılmamışsa 1)
  Uyarı           = ürün endeksi 100 ± --esik (5) dışında
  Şüpheli eşleşme = rakip fiyatı bizimkinin yarısından az ya da iki katından fazla (ambalaj farkı olabilir) — endekse girmez
Kampanyalı taramalar listelenir, endekse girmez.
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

URUN = {"kod": ("urun kodu", "stok kodu", "kod", "malzeme kodu"), "barkod": ("barkod", "ean"), "ad": ("urun adi", "urun", "aciklama", "malzeme adi"), "kategori": ("kategori", "urun grubu", "reyon"),
        "fiyat": ("fiyat", "satis fiyati", "raf fiyati", "bizim fiyat"), "adet": ("satis adedi", "satis miktari", "adet", "agirlik")}
TARAMA = {"tarih": ("tarih", "tarama tarihi"), "rakip": ("rakip", "rakip adi", "firma", "market"), "kod": ("urun kodu", "stok kodu", "kod", "bizim kod"), "barkod": ("barkod", "ean"),
          "fiyat": ("rakip fiyat", "rakip fiyati", "fiyat", "raf fiyati"), "kampanya": ("kampanya", "kampanyali", "indirim", "promosyon")}
SIFIR = Decimal(0)
EVET = {"evet", "e", "x", "var", "1", "kampanya", "kampanyali", "indirim", "indirimli"}


def oku_urun(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, URUN, ("kod", "fiyat")):
        kod = metin(r["kod"])
        if not kod:
            continue
        fiyat, adet = para(r["fiyat"]), para(r["adet"]) if metin(r["adet"]) else None
        if fiyat is None or fiyat <= 0 or (metin(r["adet"]) and (adet is None or adet < 0)):
            uy.append(uyari("Yüksek", "Ürün satırı okunamadı", kod, f"Satır {r['_satir']}: fiyat ya da satış adedi okunamadı. Ürün alınmadı."))
            continue
        if katla(kod) in gorulen:
            uy.append(uyari("Yüksek", "Ürün kodu iki kez yazılmış", kod, f"Satır {r['_satir']}: ikinci satır alınmadı."))
            continue
        gorulen.add(katla(kod))
        l.append({"kod": kod, "barkod": metin(r["barkod"]), "ad": metin(r["ad"]), "kategori": metin(r["kategori"]) or "Diğer", "fiyat": fiyat, "adet": adet})
    return l, uy


def oku_tarama(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, TARAMA, ("rakip", "fiyat")):
        rakip, kod = metin(r["rakip"]), metin(r["kod"]) or metin(r["barkod"])
        if not rakip and not kod:
            continue
        t, fiyat = tarih(r["tarih"]), para(r["fiyat"])
        if not rakip or not kod or fiyat is None or fiyat <= 0 or (metin(r["tarih"]) and t is None):
            uy.append(uyari("Yüksek", "Tarama satırı okunamadı", f"{rakip} · {kod}", f"Satır {r['_satir']}: rakip, ürün kodu, fiyat ya da tarih okunamadı. Satır alınmadı."))
            continue
        l.append({"tarih": t, "rakip": rakip, "kod": metin(r["kod"]), "barkod": metin(r["barkod"]), "fiyat": fiyat, "kampanya": katla(r["kampanya"]) in EVET, "satir": r["_satir"]})
    return l, uy


def _endeks(biz, rakip) -> Decimal | None:
    return yuvarla(biz * 100 / rakip) if rakip else None


def karsilastir(urunler: list[dict], taramalar: list[dict], esik: Decimal = Decimal(5), eski: int = 30) -> dict:
    if not urunler:
        raise ValueError("Bizim fiyat listesinde ürün yok.")
    if not taramalar:
        raise ValueError("Rakip taraması yok.")
    uy, kodla, barkodla, eslesmeyen = [], {}, {}, defaultdict(list)
    for u in urunler:
        kodla[katla(u["kod"])] = u
        if u["barkod"]:
            barkodla.setdefault(katla(u["barkod"]), u)
    son = max((t["tarih"] for t in taramalar if t["tarih"]), default=None)
    sinir = son - timedelta(days=eski) if son else None
    rakipler, gozlem = {}, {}
    for t in taramalar:
        u = kodla.get(katla(t["kod"])) or barkodla.get(katla(t["kod"])) or barkodla.get(katla(t["barkod"]))
        rk = katla(t["rakip"])
        rakipler.setdefault(rk, t["rakip"])
        if u is None:
            eslesmeyen[rk].append(t["kod"] or t["barkod"])
            continue
        g = gozlem.setdefault((katla(u["kod"]), rk), {"normal": None, "kampanya": None, "eski": None})
        alan = "kampanya" if t["kampanya"] else "eski" if sinir and t["tarih"] and t["tarih"] < sinir else "normal"
        onceki = g[alan]
        if onceki is None or (t["tarih"] or date.min, t["satir"]) >= (onceki["tarih"] or date.min, onceki["satir"]):
            g[alan] = t
    for rk, kodlar in eslesmeyen.items():
        uy.append(uyari("Bilgi", "Eşleşmeyen tarama", rakipler[rk], f"{len(kodlar)} satır bizim listedeki hiçbir ürün koduna ya da barkoda eşleşmedi: {', '.join(sorted(set(kodlar))[:8])}" + (" …" if len(set(kodlar)) > 8 else "") + "."))
    rsira = sorted(rakipler, key=lambda k: katla(rakipler[k]))
    detay, ozet = [], []
    for u in sorted(urunler, key=lambda u: (katla(u["kategori"]), katla(u["kod"]))):
        uk, fiyatlar, satir = katla(u["kod"]), {}, {"urun": u, "rakip": {}}
        for rk in rsira:
            g = gozlem.get((uk, rk))
            if not g:
                continue
            satir["rakip"][rk] = g
            if g["normal"]:
                f = g["normal"]["fiyat"]
                g["supheli"] = f * 2 < u["fiyat"] or f > u["fiyat"] * 2
                if g["supheli"]:
                    uy.append(uyari("Orta", "Eşleşme şüpheli", f"{u['kod']} · {rakipler[rk]}", f"Rakip fiyatı {f:,.2f}, bizim fiyat {u['fiyat']:,.2f}. Ambalaj ya da gramaj farkı olabilir; eşleşmeyi kontrol edin. Endekse alınmadı."))
                else:
                    fiyatlar[rk] = f
            elif g["eski"]:
                uy.append(uyari("Bilgi", "Eski tarama", f"{u['kod']} · {rakipler[rk]}", f"Son kampanyasız tarama {g['eski']['tarih']:%d.%m.%Y}; {eski} günden eski olduğu için endekse alınmadı."))
        if fiyatlar:
            ort = sum(fiyatlar.values(), SIFIR) / len(fiyatlar)
            en = min(fiyatlar.values())
            satir |= {"ort": yuvarla(ort), "en_dusuk": en, "en_ucuz": ", ".join(rakipler[k] for k in rsira if fiyatlar.get(k) == en), "endeks": _endeks(u["fiyat"], ort), "endeks_min": _endeks(u["fiyat"], en)}
            if satir["endeks"] > 100 + esik:
                uy.append(uyari("Orta", "Rakip ortalamasından pahalı", u["kod"], f"{u['ad'] or u['kod']}: bizim {u['fiyat']:,.2f}, rakip ortalaması {satir['ort']:,.2f} (endeks {satir['endeks']}; eşik {100 + esik})."))
            elif satir["endeks"] < 100 - esik:
                uy.append(uyari("Bilgi", "Rakip ortalamasından ucuz", u["kod"], f"{u['ad'] or u['kod']}: bizim {u['fiyat']:,.2f}, rakip ortalaması {satir['ort']:,.2f} (endeks {satir['endeks']}). Marj fırsatı olabilir."))
        else:
            satir |= {"ort": None, "en_dusuk": None, "en_ucuz": "", "endeks": None, "endeks_min": None}
        satir["fiyatlar"] = fiyatlar
        detay.append(satir)
    for rk in rsira:
        gruplar = defaultdict(lambda: [SIFIR, SIFIR, 0])
        for s in detay:
            if rk in s["fiyatlar"]:
                w = s["urun"]["adet"] if s["urun"]["adet"] is not None else Decimal(1)
                for kat in (s["urun"]["kategori"], None):
                    g = gruplar[katla(kat) if kat else None]
                    g[0] += s["urun"]["fiyat"] * w
                    g[1] += s["fiyatlar"][rk] * w
                    g[2] += 1
        adlar = {katla(s["urun"]["kategori"]): s["urun"]["kategori"] for s in detay}
        for k in sorted(gruplar, key=lambda k: (k is None, k or "")):
            b, r_, n = gruplar[k]
            ozet.append({"rakip": rakipler[rk], "kategori": adlar[k] if k else "TÜMÜ", "urun": n, "biz": b, "onlar": r_, "endeks": _endeks(b, r_)})
    eslesen = sum(1 for s in detay if s["fiyatlar"])
    return {"detay": detay, "ozet": ozet, "rakipler": [rakipler[k] for k in rsira], "rakip_anahtar": rsira, "eslesen": eslesen, "son": son, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, esik: Decimal) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Endeks Özeti"
    _baslik(ws, ["Rakip", "Kategori", "Karşılaştırılan Ürün", "Bizim Sepet", "Rakip Sepeti", "Fiyat Endeksi"], (20, 22, 12, 16, 16, 12))
    renk = lambda e: None if e is None else KIRMIZI if e > 100 + esik else YESIL if e < 100 - esik else None   # noqa: E731
    for o in r["ozet"]:
        k = _satir(ws, [o["rakip"], o["kategori"], o["urun"], o["biz"], o["onlar"], o["endeks"]], para_s=(4, 5, 6), dolgu={6: renk(o["endeks"])})
        if o["kategori"] == "TÜMÜ":
            for h in ws[k]:
                h.font = Font(bold=True)
    ws.append([])
    ws.append([f"Endeks = bizim ÷ rakip × 100; 100 üstü rakipten pahalı. Ağırlık: satış adedi. Kampanyalı ve eski taramalar hariç. Son tarama: {r['son']:%d.%m.%Y}." if r["son"] else "Endeks = bizim ÷ rakip × 100."])
    ws = wb.create_sheet("Ürün Karşılaştırma")
    rk = r["rakip_anahtar"]
    _baslik(ws, ["Kategori", "Ürün Kodu", "Ürün Adı", "Bizim Fiyat", "Satış Adedi"] + r["rakipler"] + ["Rakip Ortalaması", "En Düşük Rakip", "En Ucuz Rakip", "Endeks (ortalamaya)", "Endeks (en düşüğe)"],
            (16, 14, 30, 12, 10) + (12,) * len(rk) + (13, 13, 18, 12, 12), "D2")
    for s in r["detay"]:
        u = s["urun"]
        hucre = [s["fiyatlar"].get(k) for k in rk]
        n = _satir(ws, [u["kategori"], u["kod"], u["ad"], u["fiyat"], u["adet"]] + hucre + [s["ort"], s["en_dusuk"], s["en_ucuz"], s["endeks"], s["endeks_min"]],
                   para_s=[4] + [6 + i for i in range(len(rk))] + [6 + len(rk), 7 + len(rk), 9 + len(rk), 10 + len(rk)], dolgu={9 + len(rk): renk(s["endeks"])})
        for i, k in enumerate(rk):
            if k in s["fiyatlar"] and s["fiyatlar"][k] < u["fiyat"]:
                ws.cell(n, 6 + i).fill = PatternFill("solid", fgColor=SARI)
    ws = wb.create_sheet("Taramalar")
    _baslik(ws, ["Ürün Kodu", "Rakip", "Durum", "Tarih", "Fiyat"], (14, 20, 14, 12, 12))
    for s in r["detay"]:
        for k in rk:
            g = s["rakip"].get(k)
            for alan, ad in (("normal", "Endekste"), ("kampanya", "Kampanyalı"), ("eski", "Eski")):
                if g and g[alan] and not (alan == "eski" and g["normal"]):
                    ad = "Şüpheli eşleşme" if alan == "normal" and g.get("supheli") else ad
                    _satir(ws, [s["urun"]["kod"], g[alan]["rakip"], ad, g[alan]["tarih"], g[alan]["fiyat"]], para_s=(5,), tarih_s=(4,),
                           dolgu={3: MAVI if alan == "kampanya" else SARI if ad == "Şüpheli eşleşme" else None})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(urun_yol: Path, tarama_yol: Path, cikti: Path, esik: Decimal = Decimal(5), eski: int = 30) -> dict:
    urunler, uy1 = oku_urun(urun_yol)
    taramalar, uy2 = oku_tarama(tarama_yol)
    r = karsilastir(urunler, taramalar, esik, eski)
    r["uyarilar"] = sirala(uy1 + uy2 + r["uyarilar"])
    rapor_yaz(cikti, r, esik)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Rakip fiyat taramalarını ürün eşleştirmesiyle karşılaştırıp fiyat endeksi çıkarır.")
    p.add_argument("--urunler", type=Path, help="Bizim fiyatlar: Ürün Kodu, Barkod, Ürün Adı, Kategori, Fiyat, Satış Adedi")
    p.add_argument("--taramalar", type=Path, help="Rakip taramaları: Tarih, Rakip, Ürün Kodu ya da Barkod, Rakip Fiyat, Kampanya")
    p.add_argument("--esik", default="5", help="Endeks 100 ± bu değerin dışındaysa uyarılır (varsayılan 5)")
    p.add_argument("--eski", type=int, default=30, help="En son taramadan bu kadar gün eski taramalar endekse girmez (varsayılan 30)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "rakip_fiyat_karsilastirmasi.xlsx")
    a = p.parse_args(argv)
    uy, ty = a.urunler or ORNEK / "bizim_fiyatlar.csv", a.taramalar or ORNEK / "rakip_taramalari.csv"
    esik = para(a.esik)
    if esik is None or not 0 <= esik < 100 or a.eski < 0:
        print("[X] --esik 0-100 arası, --eski sıfır ya da üstü olmalı")
        return 2
    for y in (uy, ty):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(uy, ty, a.cikti, esik, a.eski)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(r['detay'])} ürün · {r['eslesen']} ürün en az bir rakiple karşılaştırıldı · {len(r['rakipler'])} rakip")
    for o in r["ozet"]:
        if o["kategori"] == "TÜMÜ":
            print(f"     {o['rakip']}: {o['urun']} ürün · endeks {o['endeks']}")
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
