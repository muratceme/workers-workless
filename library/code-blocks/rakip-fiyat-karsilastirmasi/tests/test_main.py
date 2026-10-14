import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def u(kod, kat, fiyat, adet=None, barkod=""):
    return {"kod": kod, "barkod": barkod, "ad": "", "kategori": kat, "fiyat": D(fiyat), "adet": None if adet is None else D(adet)}


def t(rakip, kod, gun, fiyat, kampanya=False, barkod="", ay=10, satir=0):
    return {"tarih": date(2026, ay, gun), "rakip": rakip, "kod": kod, "barkod": barkod, "fiyat": D(fiyat), "kampanya": kampanya, "satir": satir}


def veri():
    return ([u("P1", "K", 100, 10, "869"), u("P2", "K", 50), u("P3", "L", 20, 5)],
            [t("R1", "P1", 1, 110), t("R1", "P1", 5, 90, True), t("R1", "p1", 3, 105), t("R2", "", 5, 95, barkod="869"), t("R1", "P2", 5, 40), t("R2", "P2", 1, 60, ay=8),
             t("R1", "P3", 5, 45), t("R1", "ZZZ", 5, 1)])


class KarsilastirTesti(unittest.TestCase):
    def test_endeksler(self):
        r = main.karsilastir(*veri())
        self.assertEqual([(s["urun"]["kod"], s["ort"], s["en_dusuk"], s["en_ucuz"], s["endeks"], s["endeks_min"]) for s in r["detay"]],
                         [("P1", D("100.00"), D(95), "R2", D("100.00"), D("105.26")), ("P2", D("40.00"), D(40), "R1", D("125.00"), D("125.00")), ("P3", None, None, "", None, None)])
        self.assertEqual([(o["rakip"], o["kategori"], o["urun"], o["biz"], o["onlar"], o["endeks"]) for o in r["ozet"]],
                         [("R1", "K", 2, D(1050), D(1090), D("96.33")), ("R1", "TÜMÜ", 2, D(1050), D(1090), D("96.33")),
                          ("R2", "K", 1, D(1000), D(950), D("105.26")), ("R2", "TÜMÜ", 1, D(1000), D(950), D("105.26"))])
        self.assertEqual(sorted((x["onem"], x["tur"], x["kim"]) for x in r["uyarilar"]),
                         [("Bilgi", "Eski tarama", "P2 · R2"), ("Bilgi", "Eşleşmeyen tarama", "R1"),
                          ("Orta", "Eşleşme şüpheli", "P3 · R1"), ("Orta", "Rakip ortalamasından pahalı", "P2")])
        self.assertEqual((r["eslesen"], r["son"], r["detay"][0]["rakip"]["r1"]["kampanya"]["fiyat"]), (2, date(2026, 10, 5), D(90)))

    def test_esik_ve_eski(self):
        r = main.karsilastir(*veri(), esik=D(30))
        self.assertEqual(sorted(x["tur"] for x in r["uyarilar"] if "Rakip" in x["tur"]), [])
        r = main.karsilastir(*veri(), eski=100)
        self.assertEqual((r["detay"][1]["ort"], r["detay"][1]["endeks"], [x["tur"] for x in r["uyarilar"] if x["tur"] == "Eski tarama"]), (D("50.00"), D("100.00"), []))
        for a, b in ((([], veri()[1])), (veri()[0], [])):
            with self.assertRaises(ValueError):
                main.karsilastir(a, b)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "bizim_fiyatlar.csv", ORNEK / "rakip_taramalari.csv", cikti)
            self.assertEqual([(o["rakip"], o["urun"], o["endeks"]) for o in r["ozet"] if o["kategori"] == "TÜMÜ"], [("Rakip A", 6, D("102.20")), ("Rakip B", 6, D("98.22")), ("Rakip C", 4, D("106.58"))])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Endeks Özeti", "Ürün Karşılaştırma", "Taramalar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, z = Path(tmp) / "u.csv", Path(tmp) / "t.csv"
            y.write_text("Stok Kodu;Raf Fiyatı;Adet\nA;10,50;1.000\na;5;\nB;0;\nC;5;x\n;5;\n", encoding="utf-8")
            z.write_text("Market;Kod;Fiyat;İndirim;Tarih\nM;A;10;x;01.10.2026\nM;A;y;;\n;A;5;;\nM;;5;;\nM;A;5;;99\n;;;;\n", encoding="utf-8")
            l, uy = main.oku_urun(y)
            self.assertEqual(([(x["kod"], x["fiyat"], x["adet"], x["kategori"]) for x in l], [x["tur"] for x in uy]),
                             ([("A", D("10.50"), D(1000), "Diğer")], ["Ürün kodu iki kez yazılmış", "Ürün satırı okunamadı", "Ürün satırı okunamadı"]))
            l, uy = main.oku_tarama(z)
            self.assertEqual(([(x["rakip"], x["kod"], x["fiyat"], x["kampanya"], x["tarih"]) for x in l], len(uy)), ([("M", "A", D(10), True, date(2026, 10, 1))], 4))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--urunler", str(y), "--taramalar", str(z), "--cikti", c]
            self.assertEqual(main.main(arg + ["--esik", "2", "--eski", "7"]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--esik", "x"], ["--esik", "100"], ["--eski", "-1"])], [2] * 3)
            self.assertEqual(main.main(arg[:2] + ["--taramalar", str(Path(tmp) / "yok.csv")]), 1)
            z.write_text("Rakip;Fiyat\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
