"""
Geçici Vergi Hesaplama — Workers / Workless kod bloğu
Beyanname ve Vergi › Serbest Muhasebeci Mali Müşavir (SMMM)

Dönem sonu kümülatif mizandan ticari kârı bulur; verdiğiniz ilave, istisna, zarar ve indirimlerle geçici vergi matrahını ve ödenecek vergiyi hesaplar.
  Ticari bilanço kârı   = 600–689 arası gelir tablosu hesaplarının (alacak bakiye − borç bakiye) toplamı        (690 ve sonrası kapanış hesapları alınmaz)
  + İlaveler            = kanunen kabul edilmeyen giderler (düzeltmeler dosyasında tür: KKEG / İlave)
  − İstisnalar          = zarar olsa dahi indirilecek istisna ve indirimler (tür: İstisna); sonucu zarara çevirebilir
  − Geçmiş yıl zararı   = en çok kalan kâr kadar (tür: Geçmiş Yıl Zararı)
  − Diğer indirimler    = kazancın bulunması hâlinde indirilecekler, en çok kalan kâr kadar (tür: İndirim)
  = Matrah              (negatifse 0; dönem zararı ayrıca gösterilir)
  Hesaplanan vergi      = matrah × --oran ÷ 100
  Ödenecek              = hesaplanan − önceki dönemlerde hesaplanan geçici vergi (--onceki) − kesinti yoluyla ödenen vergi (--stopaj); negatifse 0
Kod vergi oranını ve hangi giderin kabul edilmediğini bilmez: oranı ve düzeltme kalemlerini siz verirsiniz. Beyannameden önce mutlaka bir meslek mensubu kontrol etmelidir.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


MIZAN = {"kod": ("hesap kodu", "kod", "hesap no"), "ad": ("hesap adi", "hesap", "aciklama"), "borc": ("borc bakiye", "borc bakiyesi", "bakiye borc"), "alacak": ("alacak bakiye", "alacak bakiyesi", "bakiye alacak"),
         "tborc": ("borc", "borc toplami", "borc tutari"), "talacak": ("alacak", "alacak toplami", "alacak tutari")}
DUZ = {"kalem": ("kalem", "aciklama", "duzeltme"), "tur": ("tur", "turu", "tip"), "tutar": ("tutar",)}
TURLER = {"kkeg": "ilave", "ilave": "ilave", "kanunen kabul edilmeyen gider": "ilave", "istisna": "istisna", "zarar olsa dahi indirilecek": "istisna", "gecmis yil zarari": "zarar", "gecmis yil zararlari": "zarar", "zarar": "zarar",
          "indirim": "indirim", "diger indirim": "indirim", "kazanc varsa indirilecek": "indirim"}
TUR_AD = {"ilave": "İlave (KKEG)", "istisna": "İstisna", "zarar": "Geçmiş yıl zararı", "indirim": "İndirim"}
Z = Decimal(0)


def oku_mizan(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, MIZAN, ("kod",)):
        kod = metin(r["kod"])
        if isinstance(r["kod"], float) and r["kod"] == int(r["kod"]):
            kod = str(int(r["kod"]))
        if not re.match(r"\d{3}", kod):
            continue
        b, a, tb, ta = (para(r[k]) if metin(r[k]) else None for k in ("borc", "alacak", "tborc", "talacak"))
        if b is None and a is None and (tb is not None or ta is not None):
            fark = (tb or Z) - (ta or Z)
            b, a = (fark, Z) if fark > 0 else (Z, -fark)
        if any(para(r[k]) is None for k in ("borc", "alacak", "tborc", "talacak") if metin(r[k])) or (b is None and a is None):
            uy.append(uyari("Yüksek", "Mizan satırı okunamadı", f"{kod} {metin(r['ad'])} (satır {r['_satir']})", "Bakiye tutarı okunamadı veya boş. Satır alınmadı; ticari kâr eksik hesaplanmış olabilir."))
            continue
        l.append({"kod": kod, "ad": metin(r["ad"]), "borc": b or Z, "alacak": a or Z})
    return l, uy


def oku_duzeltme(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, DUZ, ("kalem", "tur", "tutar")):
        k = metin(r["kalem"])
        if not k and not metin(r["tutar"]):
            continue
        tur, t = TURLER.get(katla(r["tur"])), para(r["tutar"])
        if tur is None or t is None or t < 0:
            uy.append(uyari("Yüksek", "Düzeltme satırı okunamadı", f"{k} (satır {r['_satir']})", f"Tür '{metin(r['tur'])}', tutar '{metin(r['tutar'])}'. Tür KKEG, İstisna, Geçmiş Yıl Zararı veya İndirim; tutar pozitif olmalı. Satır alınmadı."))
            continue
        l.append({"kalem": k, "tur": tur, "tutar": t})
    return l, uy


def ticari_kar(mizan: list[dict]) -> tuple[Decimal, list[dict], list[dict]]:
    uy = []
    gt = [h for h in mizan if "600" <= h["kod"][:3] <= "689"]
    if not gt:
        raise ValueError("Mizanda 600–689 arası gelir tablosu hesabı yok; ticari kâr hesaplanamadı (--ticari-kar ile elle verebilirsiniz).")
    ana = [h for h in gt if re.fullmatch(r"\d{3}", h["kod"])]
    if ana:
        if len(ana) < len(gt):
            uy.append(uyari("Bilgi", "Alt hesaplar toplanmadı", "Mizan", f"Gelir tablosu hesaplarında {len(gt) - len(ana)} alt hesap satırı var; mükerrer saymamak için yalnız üç haneli ana hesaplar toplandı."))
        gt = ana
    else:
        uy.append(uyari("Bilgi", "Ana hesap satırı yok", "Mizan", "Üç haneli ana hesap bulunmadı; bütün 6'lı satırlar toplandı. Dosyada hem ara toplam hem alt hesap varsa kâr mükerrer hesaplanır."))
    yedi = [h for h in mizan if h["kod"][0] == "7" and (h["borc"] or h["alacak"])]
    if yedi:
        uy.append(uyari("Yüksek", "Maliyet hesapları yansıtılmamış", "Mizan", f"7'li hesaplarda bakiye var ({', '.join(sorted({h['kod'][:3] for h in yedi})[:8])}). Yansıtma kayıtları yapılmadan gelir tablosu hesapları giderleri içermez; ticari kâr olduğundan yüksek çıkar."))
    kapanis = [h for h in mizan if "690" <= h["kod"][:3] <= "699" and (h["borc"] or h["alacak"])]
    if kapanis:
        uy.append(uyari("Bilgi", "Kapanış hesaplarında bakiye var", "Mizan", "690–699 hesapları hesaba katılmadı. Gelir tablosu hesapları 690'a devredilmişse 600–689 bakiyeleri sıfırdır; kapanış öncesi mizanı kullanın."))
    tb, ta = sum((h["borc"] for h in mizan if re.fullmatch(r"\d{3}", h["kod"]) or not ana), Z), sum((h["alacak"] for h in mizan if re.fullmatch(r"\d{3}", h["kod"]) or not ana), Z)
    if tb != ta:
        uy.append(uyari("Orta", "Mizan dengede değil", "Mizan", f"Borç bakiyeleri toplamı {tb:,.2f}, alacak bakiyeleri toplamı {ta:,.2f} (fark {tb - ta:,.2f}). Eksik veya hatalı satır olabilir."))
    return sum((h["alacak"] - h["borc"] for h in gt), Z), gt, uy


def hesapla(kar: Decimal, duzeltmeler: list[dict], oran: Decimal, onceki: Decimal = Z, stopaj: Decimal = Z) -> dict:
    t = {k: sum((d["tutar"] for d in duzeltmeler if d["tur"] == k), Z) for k in TUR_AD}
    a = kar + t["ilave"]
    b = a - t["istisna"]
    zarar_k = min(t["zarar"], max(b, Z))
    c = b - zarar_k
    indirim_k = min(t["indirim"], max(c, Z))
    d = c - indirim_k
    matrah, vergi = max(d, Z), Z
    vergi = yuvarla(matrah * oran / 100)
    kalan = vergi - onceki - stopaj
    uy = []
    if zarar_k < t["zarar"]:
        uy.append(uyari("Bilgi", "Geçmiş yıl zararının tamamı kullanılamadı", "Hesap", f"{t['zarar']:,.2f} zararın {zarar_k:,.2f} tutarı mahsup edildi; {t['zarar'] - zarar_k:,.2f} sonraki dönemlere kalır (mahsup süresini kontrol edin)."))
    if indirim_k < t["indirim"]:
        uy.append(uyari("Bilgi", "İndirimlerin tamamı kullanılamadı", "Hesap", f"{t['indirim']:,.2f} indirimin {indirim_k:,.2f} tutarı kullanıldı; kazanç yetersiz."))
    if kalan < 0:
        uy.append(uyari("Bilgi", "Mahsup edilecek tutar vergiyi aşıyor", "Hesap", f"Önceki dönem geçici vergisi ve kesintiler hesaplanan vergiden {-kalan:,.2f} fazla. Ödenecek vergi çıkmaz; fazlanın mahsup ve iade durumunu ayrıca değerlendirin."))
    return {"kar": kar, "ilave": t["ilave"], "ara1": a, "istisna": t["istisna"], "ara2": b, "zarar": t["zarar"], "zarar_k": zarar_k, "ara3": c, "indirim": t["indirim"], "indirim_k": indirim_k, "matrah": matrah,
            "donem_zarari": -d if d < 0 else Z, "oran": oran, "vergi": vergi, "onceki": onceki, "stopaj": stopaj, "odenecek": max(kalan, Z), "fazla": -kalan if kalan < 0 else Z, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, hesaplar: list[dict], duzeltmeler: list[dict], uy: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Hesap"
    _baslik(ws, ["Adım", "Tutar", "Açıklama"], (46, 18, 90))
    kalin = {"Geçici vergi matrahı", "Ödenecek geçici vergi"}
    for ad, deger, ac in (("Ticari bilanço kârı (zararı)", r["kar"], "600–689 hesaplarının alacak bakiye − borç bakiye toplamı" if hesaplar else "Elle verildi (--ticari-kar)"),
                          ("+ İlaveler (KKEG)", r["ilave"], "Düzeltmeler dosyasından"), ("= Ara toplam", r["ara1"], ""),
                          ("− İstisnalar (zarar olsa dahi indirilecek)", r["istisna"], "Sonucu zarara çevirebilir"), ("= Kâr / zarar", r["ara2"], ""),
                          ("− Geçmiş yıl zararı (kullanılan)", r["zarar_k"], f"Beyan edilen {r['zarar']:,.2f}; en çok kalan kâr kadar"), ("= İndirime esas tutar", r["ara3"], ""),
                          ("− Diğer indirimler (kullanılan)", r["indirim_k"], f"Beyan edilen {r['indirim']:,.2f}; kazanç varsa"), ("Geçici vergi matrahı", r["matrah"], f"Dönem zararı: {r['donem_zarari']:,.2f}" if r["donem_zarari"] else ""),
                          (f"Hesaplanan geçici vergi (%{r['oran'].normalize():f})", r["vergi"], "Matrah × oran"), ("− Önceki dönemlerde hesaplanan geçici vergi", r["onceki"], "--onceki"),
                          ("− Kesinti yoluyla ödenen vergi", r["stopaj"], "--stopaj"), ("Ödenecek geçici vergi", r["odenecek"], f"Mahsup fazlası: {r['fazla']:,.2f}" if r["fazla"] else "")):
        n = _satir(ws, [ad, deger, ac], para_s=(2,), sar=(3,))
        if ad in kalin:
            for c in ws[n]:
                c.font = Font(bold=True)
            ws.cell(n, 2).fill = PatternFill("solid", fgColor=SARI)
    ws.append([])
    ws.append(["Ön hesaptır. Oran, istisna ve indirimlerin kapsamı ve sırası mükellefin türüne ve yürürlükteki mevzuata göre değişir; beyannameden önce bir meslek mensubu kontrol etmelidir."])
    ws = wb.create_sheet("Gelir Tablosu Hesapları")
    _baslik(ws, ["Hesap Kodu", "Hesap Adı", "Borç Bakiye", "Alacak Bakiye", "Kâra Etkisi"], (12, 46, 18, 18, 18))
    for h in hesaplar:
        _satir(ws, [h["kod"], h["ad"], h["borc"] or None, h["alacak"] or None, h["alacak"] - h["borc"]], para_s=(3, 4, 5))
    if hesaplar:
        n = _satir(ws, ["", "Ticari bilanço kârı (zararı)", None, None, r["kar"]], para_s=(5,))
        ws.cell(n, 2).font = ws.cell(n, 5).font = Font(bold=True)
    ws = wb.create_sheet("Düzeltmeler")
    _baslik(ws, ["Kalem", "Tür", "Tutar"], (50, 22, 18))
    for d in duzeltmeler:
        _satir(ws, [d["kalem"], TUR_AD[d["tur"]], d["tutar"]], para_s=(3,))
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(mizan_yol: Path | None, cikti: Path, oran: Decimal, duzeltme_yol: Path | None = None, onceki: Decimal = Z, stopaj: Decimal = Z, kar: Decimal | None = None) -> dict:
    uy, hesaplar, duzeltmeler = [], [], []
    if kar is None:
        mizan, uy = oku_mizan(mizan_yol)
        kar, hesaplar, u2 = ticari_kar(mizan)
        uy += u2
    if duzeltme_yol:
        duzeltmeler, u3 = oku_duzeltme(duzeltme_yol)
        uy += u3
    r = hesapla(kar, duzeltmeler, oran, onceki, stopaj)
    uy += r["uyarilar"]
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, hesaplar, duzeltmeler, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Dönem mizanından geçici vergi matrahını ve ödenecek vergiyi hesaplar.")
    p.add_argument("--mizan", type=Path, default=ORNEK / "mizan.csv", help="Hesap Kodu, Hesap Adı, Borç Bakiye, Alacak Bakiye (dönem sonu kümülatif)")
    p.add_argument("--duzeltmeler", type=Path, help="Kalem, Tür (KKEG / İstisna / Geçmiş Yıl Zararı / İndirim), Tutar")
    p.add_argument("--oran", help="Geçici vergi oranı, %% (zorunlu; mükellef türüne göre siz belirlersiniz)")
    p.add_argument("--onceki", default="0", help="Önceki dönemlerde hesaplanan geçici vergi toplamı")
    p.add_argument("--stopaj", default="0", help="Yıl içinde kesinti yoluyla ödenen vergi")
    p.add_argument("--ticari-kar", help="Mizan yerine ticari bilanço kârını elle ver (zarar için negatif)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "gecici_vergi.xlsx")
    a = p.parse_args(argv)
    ornek = a.mizan == ORNEK / "mizan.csv" and a.ticari_kar is None
    duz = a.duzeltmeler or (ORNEK / "duzeltmeler.csv" if ornek else None)
    oran = para(a.oran) if a.oran else Decimal(25) if ornek else None
    onceki, stopaj = (para(a.onceki), para(a.stopaj)) if not ornek or a.oran else (Decimal(80000), Decimal(5000))
    kar = para(a.ticari_kar) if a.ticari_kar else None
    if oran is None or not 0 < oran < 100 or onceki is None or stopaj is None or onceki < 0 or stopaj < 0 or (a.ticari_kar and kar is None):
        print("[X] --oran zorunludur (0–100 arası); --onceki ve --stopaj sıfır veya pozitif, --ticari-kar sayı olmalı")
        return 2
    for y in ((a.mizan if kar is None else None), duz):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.mizan, a.cikti, oran, duz, onceki, stopaj, kar)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] Ticari kâr {r['kar']:,.2f} · ilave {r['ilave']:,.2f} · istisna {r['istisna']:,.2f} · geçmiş yıl zararı {r['zarar_k']:,.2f} · indirim {r['indirim_k']:,.2f}")
    print(f"[OK] Matrah {r['matrah']:,.2f} · hesaplanan vergi (%{oran.normalize():f}) {r['vergi']:,.2f} · ödenecek {r['odenecek']:,.2f}")
    for u in r["uyarilar"]:
        if u["onem"] != "Bilgi":
            print(f"[!] {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
