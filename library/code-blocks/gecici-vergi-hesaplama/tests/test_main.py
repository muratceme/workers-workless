import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def dz(**kw):
    return [{"kalem": k, "tur": k, "tutar": D(v)} for k, v in kw.items()]


def h(kod, borc=0, alacak=0):
    return {"kod": kod, "ad": "", "borc": D(borc), "alacak": D(alacak)}


class HesapTesti(unittest.TestCase):
    def test_tam_hesap(self):
        r = main.hesapla(D(680000), dz(ilave=45000, istisna=30000, zarar=120000, indirim=10000), D(25), D(80000), D(5000))
        self.assertEqual((r["ara1"], r["ara2"], r["zarar_k"], r["ara3"], r["indirim_k"], r["matrah"], r["donem_zarari"], r["vergi"], r["odenecek"], r["fazla"], r["uyarilar"]),
                         (D(725000), D(695000), D(120000), D(575000), D(10000), D(565000), D(0), D("141250.00"), D("56250.00"), D(0), []))

    def test_sinirlar(self):
        g = lambda kar, **kw: main.hesapla(D(kar), dz(**kw), D(20))   # noqa: E731
        r = g(100000, zarar=150000, indirim=5000)
        self.assertEqual((r["zarar_k"], r["indirim_k"], r["matrah"], r["vergi"], sorted(u["tur"] for u in r["uyarilar"])), (D(100000), D(0), D(0), D("0.00"), ["Geçmiş yıl zararının tamamı kullanılamadı", "İndirimlerin tamamı kullanılamadı"]))
        r = g(100000, istisna=130000, zarar=50000)
        self.assertEqual((r["ara2"], r["zarar_k"], r["matrah"], r["donem_zarari"], r["vergi"]), (D(-30000), D(0), D(0), D(30000), D("0.00")))
        r = g(-50000, ilave=20000)
        self.assertEqual((r["ara1"], r["matrah"], r["donem_zarari"]), (D(-30000), D(0), D(30000)))
        r = g(100000, indirim=130000)
        self.assertEqual((r["indirim_k"], r["matrah"], r["donem_zarari"]), (D(100000), D(0), D(0)))
        r = main.hesapla(D("1000.10"), [], D("22.5"), D(100), D(25))
        self.assertEqual((r["vergi"], r["odenecek"], r["fazla"]), (D("225.02"), D("100.02"), D(0)))
        r = main.hesapla(D(1000), [], D(20), D(150), D(100))
        self.assertEqual((r["vergi"], r["odenecek"], r["fazla"], [u["tur"] for u in r["uyarilar"]]), (D("200.00"), D(0), D("50.00"), ["Mahsup edilecek tutar vergiyi aşıyor"]))

    def test_ticari_kar(self):
        kar, gt, uy = main.ticari_kar([h("100", 500), h("600", 0, 1000), h("600.01", 0, 600), h("600.02", 0, 400), h("621", 700), h("689", 50), h("690", 250), h("500", 0, 500)])
        self.assertEqual((kar, [x["kod"] for x in gt], sorted(u["tur"] for u in uy)), (D(250), ["600", "621", "689"], ["Alt hesaplar toplanmadı", "Kapanış hesaplarında bakiye var"]))
        kar, gt, uy = main.ticari_kar([h("600.01", 0, 600), h("621.01", 400), h("770", 100), h("100", 100)])
        self.assertEqual((kar, sorted(u["tur"] for u in uy)), (D(200), ["Ana hesap satırı yok", "Maliyet hesapları yansıtılmamış"]))
        self.assertEqual([u["tur"] for u in main.ticari_kar([h("600", 0, 100), h("100", 90)])[2]], ["Mizan dengede değil"])
        with self.assertRaises(ValueError):
            main.ticari_kar([h("100", 5), h("500", 0, 5)])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "g.xlsx"
            r = main.calistir(ORNEK / "mizan.csv", cikti, D(25), ORNEK / "duzeltmeler.csv", D(80000), D(5000))
            self.assertEqual((r["kar"], r["matrah"], r["vergi"], r["odenecek"], r["uyarilar"]), (D("680000.00"), D("565000.00"), D("141250.00"), D("56250.00"), []))
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Hesap", "Gelir Tablosu Hesapları", "Düzeltmeler", "Uyarılar"])
            self.assertEqual((wb["Gelir Tablosu Hesapları"].max_row, wb["Hesap"]["B10"].value), (9, 565000))

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            my, dy = Path(tmp) / "m.csv", Path(tmp) / "d.csv"
            my.write_text("Kod;Hesap;Borç;Alacak\n600;Satış;100;1.100,50\n621;Maliyet;700;\nTOPLAM;;800;1100,50\n632;Gider;x;\n100;Kasa;300,50;\n", encoding="utf-8")
            dy.write_text("Açıklama;Tür;Tutar\nCeza;KKEG;50\nBağış;Bilinmiyor;5\nZarar;Geçmiş Yıl Zararı;-5\n", encoding="utf-8")
            m, uy = main.oku_mizan(my)
            self.assertEqual(([(x["kod"], x["borc"], x["alacak"]) for x in m], [u["tur"] for u in uy]), ([("600", D(0), D("1000.50")), ("621", D(700), D(0)), ("100", D("300.50"), D(0))], ["Mizan satırı okunamadı"]))
            d, uy = main.oku_duzeltme(dy)
            self.assertEqual((d, [u["tur"] for u in uy]), ([{"kalem": "Ceza", "tur": "ilave", "tutar": D(50)}], ["Düzeltme satırı okunamadı"] * 2))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--mizan", str(my), "--cikti", c]
            self.assertEqual(main.main(arg + ["--oran", "25", "--duzeltmeler", str(dy), "--onceki", "10", "--stopaj", "5"]), 0)
            self.assertEqual(main.main(["--ticari-kar=-5.000,00", "--oran", "20", "--cikti", c]), 0)
            for ek in ([], ["--oran", "0"], ["--oran", "25", "--onceki", "-1"], ["--oran", "25", "--ticari-kar", "x"]):
                self.assertEqual(main.main(arg + ek), 2)
            self.assertEqual(main.main(["--mizan", str(Path(tmp) / "yok.csv"), "--oran", "25"]), 1)
            my.write_text("Kod;Hesap;Borç Bakiye;Alacak Bakiye\n100;Kasa;5;\n500;Sermaye;;5\n", encoding="utf-8")
            self.assertEqual(main.main(arg + ["--oran", "25"]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
