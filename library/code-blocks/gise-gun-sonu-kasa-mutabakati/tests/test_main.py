import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def sayim(**kw):
    """sayim(G1_TL={200: 10, 50: 4}) → {("G1", "TL"): {200: 10, 50: 4}}"""
    return {tuple(k.split("_")): {D(str(a)): D(b) for a, b in v.items()} for k, v in kw.items()}


def sistem(**kw):
    return {tuple(k.split("_")): {"sistem": D(str(v[0])), "devir": D(str(v[1])) if len(v) > 1 else None, "limit": D(str(v[2])) if len(v) > 2 else None} for k, v in kw.items()}


def islem(no, gise, pb, giris=0, cikis=0, kur=None):
    return {"no": no, "saat": "10:00", "gise": gise, "tur": "İşlem", "pb": pb, "giris": D(str(giris)), "cikis": D(str(cikis)), "kur": D(str(kur)) if kur else None, "aciklama": ""}


def ipucu(r):
    return [(i["tur"], i["no"]) for i in r["ipuclari"]]


class MutabakatTesti(unittest.TestCase):
    def test_mutabik(self):
        r = main.mutabakat(sayim(G1_TL={200: 10, 50: 4, "0.50": 3}), sistem(G1_TL=("2201.50", 2000)), [islem("F1", "G1", "TL", giris="201.50")])
        s = r["satirlar"][0]
        self.assertEqual((s["sayim"], s["hesap"], s["fark"], s["durum"]), (D("2201.50"), D("2201.50"), D("0.00"), "Mutabık"))
        self.assertEqual(r["ipuclari"], [])
        self.assertEqual(r["uyarilar"], [])

    def test_tolerans_ve_yon(self):
        s = sayim(G1_TL={100: 10})
        self.assertEqual(main.mutabakat(s, sistem(G1_TL=("1000.50",)), [])["satirlar"][0]["durum"], "Kasa noksanı")
        self.assertEqual(main.mutabakat(s, sistem(G1_TL=("1000.50",)), [], D(1))["satirlar"][0]["durum"], "Mutabık")
        self.assertEqual(main.mutabakat(s, sistem(G1_TL=("999",)), [])["satirlar"][0]["durum"], "Kasa fazlası")

    def test_sistem_islemlerle_tutmuyor_ve_limit(self):
        r = main.mutabakat(sayim(G1_TL={100: 12}), sistem(G1_TL=(1200, 1000, 1100)), [islem("F1", "G1", "TL", giris=150)])
        self.assertEqual(sorted(u["tur"] for u in r["uyarilar"]), ["Kasa limiti aşıldı", "Sistem bakiyesi işlemlerle tutmuyor"])

    def test_eksik_taraf(self):
        r = main.mutabakat(sayim(G1_TL={100: 1}), sistem(G2_TL=(500,)), [])
        self.assertEqual([s["durum"] for s in r["satirlar"]], ["Sistem yok", "Kasa noksanı"])
        self.assertIn("Sayım yok", [u["tur"] for u in r["uyarilar"]])


class IpucuTesti(unittest.TestCase):
    def test_tutar_esit(self):
        isl = [islem("F1", "G1", "TL", giris=700), islem("F2", "G1", "TL", cikis=700), islem("F3", "G1", "TL", giris=300)]
        r = main.mutabakat(sayim(G1_TL={100: 3}), sistem(G1_TL=(1000,)), isl)                  # noksan 700
        self.assertEqual(ipucu(r), [("Tutar farka eşit", "F1"), ("Tutar farka eşit", "F2")])
        self.assertIn("girmemiş", r["ipuclari"][0]["aciklama"])                                 # giriş kaydı var, para yok
        self.assertIn("iki kez ödenmiş", r["ipuclari"][1]["aciklama"])

    def test_ters_yon(self):
        isl = [islem("F1", "G1", "TL", cikis=1000), islem("F2", "G1", "TL", giris=1000)]
        r = main.mutabakat(sayim(G1_TL={200: 15}), sistem(G1_TL=(1000,)), isl)                 # fazla 2.000
        self.assertEqual([x for x in ipucu(r) if x[0] == "Ters yön kaydı"], [("Ters yön kaydı", "F1")])   # yalnız çıkış kaydı fazla verir

    def test_rakam_yer_degisimi(self):
        self.assertEqual(main.yer_degisimi(D("5400"), D("900")), "4500.00")
        self.assertEqual(main.yer_degisimi(D("1234.56"), D("0.09")), "1234.65")
        self.assertIsNone(main.yer_degisimi(D("5500"), D("900")))
        r = main.mutabakat(sayim(G1_TL={100: 91}), sistem(G1_TL=(10000,)), [islem("F1", "G1", "TL", giris=5400), islem("F2", "G1", "TL", giris=2200)])
        self.assertEqual(ipucu(r), [("Rakam yer değişimi", "F1")])

    def test_kupur(self):
        r = main.mutabakat(sayim(G1_USD={100: 8, 20: 5}), sistem(G1_USD=(1100,)), [])          # noksan 200 = 2 × 100
        self.assertEqual(ipucu(r), [("Kupür sayımı", "")])
        self.assertIn("2 adet 100 USD", r["ipuclari"][0]["aciklama"])

    def test_giseler_arasi_ve_doviz(self):
        r = main.mutabakat(sayim(G1_TL={200: 7}, G2_TL={200: 3}), sistem(G1_TL=(1000,), G2_TL=(1000,)), [])    # +400 ve −400
        self.assertEqual(sorted((i["gise"], i["tur"]) for i in r["ipuclari"] if i["tur"] == "Gişeler arası"), [("G1", "Gişeler arası"), ("G2", "Gişeler arası")])
        isl = [islem("F9", "G1", "USD", cikis=1000, kur="40.50")]
        r = main.mutabakat(sayim(G1_TL={50: 405}, G1_USD={100: 5}), sistem(G1_TL=(0,), G1_USD=(1000,)), isl)   # TL +20.250, USD −500 → 40,5
        d = [i for i in r["ipuclari"] if i["tur"] == "Döviz işlemi"]
        self.assertEqual((len(d), d[0]["no"]), (1, "F9"))
        self.assertIn("40.5000", d[0]["aciklama"])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "k.xlsx"
            r = main.calistir(ORNEK / "sayim.csv", ORNEK / "sistem.csv", cikti, ORNEK / "islemler.csv")
            self.assertEqual({(s["gise"], s["pb"]): s["fark"] for s in r["satirlar"] if s["durum"] != "Mutabık"},
                             {("G2", "TL"): D(-900), ("G3", "TL"): D(2000), ("G4", "TL"): D(20250), ("G4", "USD"): D(-500)})
            self.assertTrue(all(s["hesap"] == s["sistem"] for s in r["satirlar"]))
            self.assertIn(("Rakam yer değişimi", "F26140"), ipucu(r))
            self.assertIn(("Ters yön kaydı", "F26142"), ipucu(r))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Mutabakat", "Kupür Dökümü", "Araştırma Listesi", "Uyarılar"])

    def test_okuma(self):
        with tempfile.TemporaryDirectory() as tmp:
            s, b = Path(tmp) / "s.csv", Path(tmp) / "b.csv"
            s.write_text("Gişe;Para Birimi;Kupür;Adet\nG1;TRY;200;5\nG1;TL;0,50;4\nG1;TL;75;1\nG1;TL;100;iki\n", encoding="utf-8")
            b.write_text("Gişe;Para Birimi;Sistem Bakiyesi\nG1;TL;1.077,00\n", encoding="utf-8")
            sy, ss, isl, uy = main.oku(s, b, None)
            self.assertEqual(sum(k * a for k, a in sy[("G1", "TL")].items()), D("1077.00"))    # TRY = TL; 75 uyarıyla toplama katılır
            self.assertEqual(sorted(u["tur"] for u in uy), ["Sayım satırı okunamadı", "Tanınmayan kupür"])
            self.assertEqual(main.main(["--sayim", str(s), "--sistem", str(b), "--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--sayim", str(Path(tmp) / "yok.csv")]), 1)
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "y.xlsx")]), 0)


if __name__ == "__main__":
    unittest.main()
