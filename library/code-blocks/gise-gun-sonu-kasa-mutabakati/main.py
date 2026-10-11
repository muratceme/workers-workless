"""
Gişe Gün Sonu Kasa Mutabakatı — Workers / Workless kod bloğu
Şube Bankacılığı › Gişe Yetkilisi

Gün sonu kasa sayımını (kupür bazında) sistem bakiyesiyle gişe ve para birimi bazında karşılaştırır; fark varsa günün
işlemleri içinden araştırılacak adayları listeler.
  Sayım          = Σ kupür × adet
  Hesaplanan     = devir bakiyesi + günün girişleri − çıkışları   (işlem dosyası verilirse)
  Fark           = sayım − sistem bakiyesi   (+ kasa fazlası, − kasa noksanı)
Araştırma ipuçları (her biri bir olasılıktır, kanıt değildir):
  - Tutarı farka eşit işlem          → kaydedilip para hareketi yapılmamış veya iki kez ödenmiş / alınmış olabilir.
  - Tutarı farkın yarısı olan işlem  → giriş yerine çıkış (veya tersi) kaydedilmiş olabilir.
  - Fark 9'a bölünüyorsa             → rakam yer değişimi (ör. 5.400 yerine 4.500); iki komşu rakamı yer değiştirince
                                       fark kadar değişen işlemler listelenir.
  - Fark tek bir kupürün katıysa     → o kupürden eksik / fazla sayım veya ödeme.
  - Aynı para biriminde bir gişe fazla, diğeri aynı tutarda noksansa → gişeler arası devir kaydı.
  - Aynı gişede TL ve döviz farkı ters yönlüyse → döviz alım-satım işlemi; fark oranı işlemdeki kurla karşılaştırılır.
Kasa limiti girilmişse limit aşımı da raporlanır. İnternete bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: 4 gişe, TL / USD / EUR
    python main.py --sayim sayim.xlsx --sistem sistem.xlsx --islemler islemler.xlsx
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SAYIM_SUTUNLARI = {"gise": ("gise", "vezne", "kasa", "gise no"), "pb": ("para birimi", "doviz cinsi", "pb"), "kupur": ("kupur", "kupur degeri"), "adet": ("adet", "sayi")}
SISTEM_SUTUNLARI = {"gise": ("gise", "vezne", "kasa", "gise no"), "pb": ("para birimi", "doviz cinsi", "pb"), "devir": ("devir bakiyesi", "devir", "acilis bakiyesi"),
                    "sistem": ("sistem bakiyesi", "gun sonu bakiyesi", "bakiye"), "limit": ("kasa limiti", "limit")}
ISLEM_SUTUNLARI = {"no": ("islem no", "fis no", "referans"), "saat": ("saat", "islem saati"), "gise": ("gise", "vezne", "kasa", "gise no"), "tur": ("islem turu", "islem", "tur"),
                   "pb": ("para birimi", "doviz cinsi", "pb"), "giris": ("giris", "kasa giris", "tahsilat"), "cikis": ("cikis", "kasa cikis", "odeme"),
                   "kur": ("kur",), "aciklama": ("aciklama", "musteri")}
# Tedavüldeki kupürler (bilinmeyen kupür uyarı verir, toplamdan düşülmez)
KUPURLER = {"TL": ("200", "100", "50", "20", "10", "5", "1", "0.50", "0.25", "0.10", "0.05", "0.01"), "USD": ("100", "50", "20", "10", "5", "2", "1"),
            "EUR": ("500", "200", "100", "50", "20", "10", "5", "2", "1", "0.50", "0.20", "0.10", "0.05", "0.02", "0.01"), "GBP": ("50", "20", "10", "5", "2", "1"),
            "CHF": ("1000", "200", "100", "50", "20", "10", "5", "2", "1", "0.50")}


def pb_duzelt(x) -> str:
    k = metin(x).upper().replace("TRY", "TL").replace("₺", "TL")
    return k or "TL"


def oku(sayim_yol: Path, sistem_yol: Path, islem_yol: Path | None) -> tuple[dict, dict, list[dict], list[dict]]:
    uy, sayim, sistem, islemler = [], defaultdict(lambda: defaultdict(Decimal)), {}, []
    for r in kayitlar(sayim_yol, SAYIM_SUTUNLARI, ("gise", "kupur", "adet")):
        if not metin(r["gise"]):
            continue
        k, a, pb = para(r["kupur"]), para(r["adet"]), pb_duzelt(r["pb"])
        if k is None or a is None or k <= 0 or a < 0 or a != a.to_integral_value():
            uy.append(uyari("Yüksek", "Sayım satırı okunamadı", r["gise"], f"Satır {r['_satir']}: kupür '{metin(r['kupur'])}', adet '{metin(r['adet'])}'. Satır toplama katılmadı."))
            continue
        if pb in KUPURLER and k not in {Decimal(x) for x in KUPURLER[pb]}:
            uy.append(uyari("Orta", "Tanınmayan kupür", r["gise"], f"Satır {r['_satir']}: {k:f} {pb} tedavüldeki kupürler arasında yok. Toplama katıldı; kontrol edin."))
        sayim[(metin(r["gise"]), pb)][k] += a
    for r in kayitlar(sistem_yol, SISTEM_SUTUNLARI, ("gise", "sistem")):
        if not metin(r["gise"]):
            continue
        s = para(r["sistem"])
        if s is None:
            uy.append(uyari("Yüksek", "Sistem bakiyesi okunamadı", r["gise"], f"Satır {r['_satir']}: '{metin(r['sistem'])}'."))
            continue
        sistem[(metin(r["gise"]), pb_duzelt(r["pb"]))] = {"sistem": s, "devir": para(r["devir"]), "limit": para(r["limit"])}
    if islem_yol:
        for r in kayitlar(islem_yol, ISLEM_SUTUNLARI, ("gise",)):
            g, c = para(r["giris"]) or Decimal(0), para(r["cikis"]) or Decimal(0)
            if metin(r["gise"]) and (g or c):
                islemler.append({"no": metin(r["no"]), "saat": metin(r["saat"])[:5], "gise": metin(r["gise"]), "tur": metin(r["tur"]), "pb": pb_duzelt(r["pb"]),
                                 "giris": g, "cikis": c, "kur": para(r["kur"]), "aciklama": metin(r["aciklama"])})
    return sayim, sistem, islemler, uy


def yer_degisimi(tutar: Decimal, fark: Decimal) -> str | None:
    """Tutarın iki komşu rakamı yer değiştirince |fark| kadar değişiyorsa değişmiş hâlini döndürür."""
    s = f"{tutar:.2f}"
    rakam = [i for i, c in enumerate(s) if c.isdigit()]
    for a, b in zip(rakam, rakam[1:]):
        if s[a] != s[b]:
            l = list(s)
            l[a], l[b] = l[b], l[a]
            if abs(Decimal("".join(l)) - tutar) == abs(fark):
                return "".join(l)
    return None


def tl(x: Decimal) -> str:
    return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def mutabakat(sayim: dict, sistem: dict, islemler: list[dict], tolerans: Decimal = Decimal(0)) -> dict:
    uy, satirlar, ipuclari = [], [], []
    for anahtar in sorted(set(sayim) | set(sistem)):
        gise, pb = anahtar
        sy = sum((k * a for k, a in sayim.get(anahtar, {}).items()), Decimal(0))
        ss = sistem.get(anahtar)
        isl = [i for i in islemler if (i["gise"], i["pb"]) == anahtar]
        giris, cikis = sum((i["giris"] for i in isl), Decimal(0)), sum((i["cikis"] for i in isl), Decimal(0))
        if ss is None:
            uy.append(uyari("Yüksek", "Sistem bakiyesi yok", f"{gise} {pb}", f"Sayımda {tl(sy)} {pb} var, sistem dosyasında bu gişe ve para birimi yok."))
            satirlar.append({"gise": gise, "pb": pb, "devir": None, "giris": giris, "cikis": cikis, "hesap": None, "sistem": None, "sayim": sy, "fark": None, "durum": "Sistem yok"})
            continue
        if anahtar not in sayim:
            uy.append(uyari("Orta" if ss["sistem"] else "Bilgi", "Sayım yok", f"{gise} {pb}", f"Sistem bakiyesi {tl(ss['sistem'])} {pb}, sayım girilmemiş (sıfır sayıldı)."))
        hesap = ss["devir"] + giris - cikis if ss["devir"] is not None and islemler else None
        if hesap is not None and hesap != ss["sistem"]:
            uy.append(uyari("Yüksek", "Sistem bakiyesi işlemlerle tutmuyor", f"{gise} {pb}",
                            f"Devir {tl(ss['devir'])} + giriş {tl(giris)} − çıkış {tl(cikis)} = {tl(hesap)}; sistem bakiyesi {tl(ss['sistem'])}. İşlem dökümü eksik olabilir."))
        fark = sy - ss["sistem"]
        durum = "Mutabık" if abs(fark) <= tolerans else "Kasa fazlası" if fark > 0 else "Kasa noksanı"
        if ss["limit"] is not None and sy > ss["limit"]:
            uy.append(uyari("Orta", "Kasa limiti aşıldı", f"{gise} {pb}", f"Sayım {tl(sy)} {pb}, limit {tl(ss['limit'])}. Fazla {tl(sy - ss['limit'])} ana kasaya devredilmeli."))
        satirlar.append({"gise": gise, "pb": pb, "devir": ss["devir"], "giris": giris, "cikis": cikis, "hesap": hesap, "sistem": ss["sistem"], "sayim": sy, "fark": fark,
                         "durum": durum, "limit": ss["limit"]})
        if durum == "Mutabık":
            continue
        uy.append(uyari("Yüksek", durum, f"{gise} {pb}", f"Sayım {tl(sy)}, sistem {tl(ss['sistem'])}, fark {'+' if fark > 0 else ''}{tl(fark)} {pb}."))
        f = abs(fark)

        def ekle(tur, aciklama, i=None):
            ipuclari.append({"gise": gise, "pb": pb, "fark": fark, "tur": tur, "aciklama": aciklama, "no": i["no"] if i else "", "saat": i["saat"] if i else "",
                             "islem": i["tur"] if i else "", "tutar": (i["giris"] or -i["cikis"]) if i else None, "not": i["aciklama"] if i else ""})
        for i in isl:
            t, yon = (i["giris"], "giriş") if i["giris"] else (i["cikis"], "çıkış")
            if t == f:
                if (fark < 0) == (yon == "giriş"):
                    ekle("Tutar farka eşit", f"{yon.capitalize()} kaydı var; para {'kasaya girmemiş' if yon == 'giriş' else 'kasadan çıkmamış'} olabilir (fiş ve müşteri imzasını kontrol edin).", i)
                else:
                    ekle("Tutar farka eşit", f"{yon.capitalize()} iki kez {'alınmış' if yon == 'giriş' else 'ödenmiş'} veya kaydı eksik bir işlem olabilir.", i)
            if t * 2 == f and (fark > 0) == (yon == "çıkış"):
                ekle("Ters yön kaydı", f"İşlem {yon} olarak kaydedilmiş; gerçekte {'giriş' if yon == 'çıkış' else 'çıkış'} ise fark tam bu tutarın iki katıdır.", i)
            yd = yer_degisimi(t, f) if (f * 100) % 9 == 0 else None
            if yd:
                ekle("Rakam yer değişimi", f"Kayıtlı tutar {tl(t)}; rakamlar yer değiştirince {tl(Decimal(yd))} olur, fark {tl(f)}. Fişteki tutarla karşılaştırın.", i)
        for k in sorted(sayim.get(anahtar, {}), reverse=True):
            if k >= 1 and f % k == 0 and f / k <= 5:
                ekle("Kupür sayımı", f"Fark {int(f / k)} adet {k:f} {pb} kupüre eşit. Bu kupürü yeniden sayın; {'eksik' if fark < 0 else 'fazla'} ödeme / tahsilat olabilir.")
                break
    farklar = [s for s in satirlar if s["durum"] in ("Kasa fazlası", "Kasa noksanı")]
    for a in farklar:
        for b in farklar:
            if a["gise"] < b["gise"] and a["pb"] == b["pb"] and a["fark"] == -b["fark"]:
                for x, y in ((a, b), (b, a)):
                    ipuclari.append({"gise": x["gise"], "pb": x["pb"], "fark": x["fark"], "tur": "Gişeler arası", "no": "", "saat": "", "islem": "", "tutar": None, "not": "",
                                     "aciklama": f"{y['gise']} gişesinde aynı tutarda ters fark var; gişeler arası devir veya para alışverişi kayda geçmemiş olabilir."})
            if a["gise"] == b["gise"] and a["pb"] == "TL" and b["pb"] != "TL" and (a["fark"] > 0) != (b["fark"] > 0):
                oran = yuvarla(abs(a["fark"] / b["fark"]), "0.0001")
                kambiyo = [i for i in islemler if i["gise"] == a["gise"] and i["pb"] == b["pb"] and i["kur"] and abs(i["kur"] - oran) / i["kur"] <= Decimal("0.02")]
                for i in kambiyo or [None]:
                    ipuclari.append({"gise": a["gise"], "pb": f"TL / {b['pb']}", "fark": a["fark"], "tur": "Döviz işlemi", "no": i["no"] if i else "", "saat": i["saat"] if i else "",
                                     "islem": i["tur"] if i else "", "tutar": (i["giris"] or -i["cikis"]) if i else None, "not": i["aciklama"] if i else "",
                                     "aciklama": f"TL ve {b['pb']} farkları ters yönlü; oran {oran:f}. " + ("İşlem kuru bu orana yakın; TL veya döviz bacağı eksik / ters olabilir." if i
                                                                                                         else "Günün döviz alım-satım işlemlerini kontrol edin.")})
    return {"satirlar": satirlar, "ipuclari": ipuclari, "uyarilar": uy, "sayim": sayim}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

DURUM_RENK = {"Mutabık": YESIL, "Kasa fazlası": SARI, "Kasa noksanı": KIRMIZI, "Sistem yok": KIRMIZI}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    mu = wb.active
    mu.title = "Mutabakat"
    _baslik(mu, ["Gişe", "Para Birimi", "Devir", "Giriş", "Çıkış", "Hesaplanan", "Sistem Bakiyesi", "Sayım", "Fark", "Durum", "Kasa Limiti", "Sayan", "Kontrol Eden"],
            (12, 9, 15, 15, 15, 15, 15, 15, 13, 14, 14, 16, 16))
    for s in r["satirlar"]:
        _satir(mu, [s["gise"], s["pb"], s["devir"], s["giris"], s["cikis"], s["hesap"], s["sistem"], s["sayim"], s["fark"], s["durum"], s.get("limit"), "", ""],
               para_s=(3, 4, 5, 6, 7, 8, 9, 11), dolgu={10: DURUM_RENK[s["durum"]], 12: SARI, 13: SARI})
    ku = wb.create_sheet("Kupür Dökümü")
    _baslik(ku, ["Gişe", "Para Birimi", "Kupür", "Adet", "Tutar"], (12, 9, 10, 9, 15))
    for (gise, pb), kupurler in sorted(r["sayim"].items()):
        for k in sorted(kupurler, reverse=True):
            _satir(ku, [gise, pb, k, int(kupurler[k]), k * kupurler[k]], para_s=(3, 5))
    ar = wb.create_sheet("Araştırma Listesi")
    _baslik(ar, ["Gişe", "Para Birimi", "Fark", "İpucu", "İşlem No", "Saat", "İşlem Türü", "Tutar (+giriş / −çıkış)", "İşlem Açıklaması", "Neye Bakılmalı", "Sonuç"],
            (12, 10, 13, 20, 12, 7, 20, 15, 28, 90, 24))
    for i in r["ipuclari"]:
        _satir(ar, [i["gise"], i["pb"], i["fark"], i["tur"], i["no"], i["saat"], i["islem"], i["tutar"], i["not"], i["aciklama"], ""], para_s=(3, 8), dolgu={11: SARI}, sar=(10,))
    if not r["ipuclari"]:
        ar.append(["Fark yok veya farkı açıklayan aday işlem bulunamadı."])
    ar.append([])
    ar.append(["İpuçları olasılıktır; fark ancak fiş, dekont ve kamera kaydı incelenerek açıklanır."])
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(sayim_yol: Path, sistem_yol: Path, cikti: Path, islem_yol: Path | None = None, tolerans: Decimal = Decimal(0)) -> dict:
    sayim, sistem, islemler, uy = oku(sayim_yol, sistem_yol, islem_yol)
    if not sayim and not sistem:
        raise ValueError("Sayım ve sistem dosyalarında okunabilir satır yok")
    r = mutabakat(sayim, sistem, islemler, tolerans)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Gün sonu kasa sayımını sistem bakiyesiyle karşılaştırır, farklar için araştırma listesi çıkarır.")
    p.add_argument("--sayim", type=Path, default=ORNEK / "sayim.csv", help="Gişe, Para Birimi, Kupür, Adet")
    p.add_argument("--sistem", type=Path, default=ORNEK / "sistem.csv", help="Gişe, Para Birimi, Devir Bakiyesi, Sistem Bakiyesi, Kasa Limiti")
    p.add_argument("--islemler", type=Path, help="İşlem No, Saat, Gişe, İşlem Türü, Para Birimi, Giriş, Çıkış, Kur, Açıklama")
    p.add_argument("--tolerans", type=Decimal, default=Decimal(0), help="Bu tutara kadar fark mutabık sayılır (varsayılan 0)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kasa_mutabakati.xlsx")
    a = p.parse_args(argv)
    isl = a.islemler or (ORNEK / "islemler.csv" if a.sayim == ORNEK / "sayim.csv" else None)
    for y in (a.sayim, a.sistem, isl):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.sayim, a.sistem, a.cikti, isl, a.tolerans)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    farkli = [s for s in r["satirlar"] if s["durum"] != "Mutabık"]
    print(f"[OK] {len(r['satirlar'])} gişe / para birimi · {len(r['satirlar']) - len(farkli)} mutabık · {len(farkli)} farklı · {len(r['ipuclari'])} araştırma ipucu")
    for s in farkli:
        print(f"[!] {s['gise']} {s['pb']}: {s['durum']}" + (f" {tl(s['fark'])}" if s["fark"] is not None else ""))
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
