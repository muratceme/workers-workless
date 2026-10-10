"""
İhracat Evrak Kontrol Listesi — Workers / Workless kod bloğu
Dış Ticaret › Dış Ticaret Uzmanı

Ülke, teslim şekli, taşıma ve ödeme yöntemine göre gereken ihracat evraklarını listeler ve hazırlanan evraklardaki
bilgilerin tutarlılığını kontrol eder:
  - Gereken evraklar bir kural tablosundan çıkar (koşul: her zaman / ülke grubu / taşıma / teslim şekli / ödeme / ürün
    özelliği). Varsayılan tablo kodda tanımlıdır; --kurallar ile kendi tablonuz eklenir veya aynı adlı kural değiştirilir.
  - Ülke grubu: AB üyesi ülkeler kodda tanımlıdır; STA ülkeleri için sevkiyat bilgisine "Ülke Grubu;STA" yazılır.
  - Dolaşım / menşe belgesi: AB'ye sanayi ve işlenmiş tarım ürünü → A.TR; AB'ye tarım veya kömür-çelik ürünü ve STA
    ülkeleri → EUR.1 / EUR-MED (veya fatura beyanı); diğer ülkeler → Menşe Şahadetnamesi (alıcı / ülke istiyorsa).
  - Hazırlanan evraklar (Evrak; Alan; Değer) verilirse: eksik zorunlu evrak, ve aynı alanın (alıcı, tutar, döviz, koli,
    net / brüt ağırlık, GTİP, menşe, teslim şekli...) evraklar arasında farklı yazılması.
Rapor: kontrol listesi (zorunluluk, gerekçe, düzenleyen, durum), tutarlılık matrisi, uyarılar. İnternete bağlanmaz.

Kullanım:
    python main.py                                                  # örnek: Almanya, CIF, deniz, vesaik mukabili
    python main.py --sevkiyat sevkiyat.csv --evraklar evraklar.xlsx --kurallar ek_kurallar.csv
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")

AB_ULKELERI = {"almanya", "avusturya", "belcika", "bulgaristan", "cekya", "cek cumhuriyeti", "danimarka", "estonya", "finlandiya", "fransa", "hirvatistan",
               "hollanda", "irlanda", "ispanya", "isvec", "italya", "kibris", "gkry", "letonya", "litvanya", "luksemburg", "macaristan", "malta", "polonya",
               "portekiz", "romanya", "slovakya", "slovenya", "yunanistan",
               "germany", "austria", "belgium", "bulgaria", "czechia", "czech republic", "denmark", "estonia", "finland", "france", "croatia", "netherlands",
               "ireland", "spain", "sweden", "italy", "cyprus", "latvia", "lithuania", "luxembourg", "hungary", "poland", "portugal", "romania", "slovakia",
               "slovenia", "greece"}

# (evrak, koşul, zorunluluk, düzenleyen, açıklama). Koşul: "her zaman" veya "alan=değer1|değer2" ; birden çok koşul " & " ile.
VARSAYILAN_KURALLAR = [
    ("Ticari Fatura (Commercial Invoice)", "her zaman", "Zorunlu", "İhracatçı", "e-Fatura / e-Arşiv ihracat faturası; alıcı ülke diline veya İngilizceye uygun nüsha"),
    ("Çeki Listesi (Packing List)", "her zaman", "Zorunlu", "İhracatçı", "Koli, net ve brüt ağırlık, hacim"),
    ("İhracat Gümrük Beyannamesi", "her zaman", "Zorunlu", "Gümrük müşaviri", "Gümrük idaresine elektronik beyan"),
    ("İhracatçı Birliği kaydı / onayı", "her zaman", "Zorunlu", "İhracatçı / birlik", "İhracatçı birliği üyeliği ve beyanname onayı"),
    ("Konşimento (Bill of Lading)", "tasima=deniz", "Zorunlu", "Taşıyıcı / acente", "Deniz taşıma senedi"),
    ("CMR Taşıma Senedi", "tasima=karayolu", "Zorunlu", "Taşıyıcı", "Karayolu taşıma senedi"),
    ("Hava Yük Senedi (AWB)", "tasima=havayolu", "Zorunlu", "Taşıyıcı / acente", "Havayolu taşıma senedi"),
    ("CIM Taşıma Senedi", "tasima=demiryolu", "Zorunlu", "Taşıyıcı", "Demiryolu taşıma senedi"),
    ("A.TR Dolaşım Belgesi", "ulke grubu=ab & urun turu=sanayi|islenmis tarim", "Zorunlu", "İhracatçı / oda onayı, gümrük vizesi",
     "Türkiye–AB Gümrük Birliği kapsamındaki ürünlerde serbest dolaşımı gösterir"),
    ("EUR.1 / EUR-MED Dolaşım Sertifikası veya fatura beyanı", "ulke grubu=ab & urun turu=tarim|komur celik", "Zorunlu", "İhracatçı / oda onayı, gümrük vizesi",
     "AB'ye tarım ve kömür-çelik ürünlerinde tercihli menşe ispatı"),
    ("EUR.1 / EUR-MED Dolaşım Sertifikası veya fatura beyanı", "ulke grubu=sta", "Zorunlu", "İhracatçı / oda onayı, gümrük vizesi",
     "Serbest ticaret anlaşması kapsamında tercihli menşe ispatı; anlaşmanın menşe kurallarını kontrol edin"),
    ("Menşe Şahadetnamesi (Certificate of Origin)", "ulke grubu=diger", "Önerilir", "İhracatçı / ticaret veya sanayi odası",
     "Alıcı, alıcı ülke gümrüğü veya akreditif istiyorsa zorunludur"),
    ("Sigorta Poliçesi / Sertifikası", "teslim=cif|cip", "Zorunlu", "Sigorta şirketi", "CIF ve CIP teslimde satıcı sigorta yaptırır (mal bedelinin en az %110'u)"),
    ("Akreditifte istenen belgeler (46A alanı)", "odeme=akreditif", "Zorunlu", "İhracatçı", "Akreditif metnindeki belge listesi, nüsha sayısı ve ibraz süresi esas alınır"),
    ("Poliçe (Draft / Bill of Exchange)", "odeme=akreditif|vesaik mukabili", "Önerilir", "İhracatçı", "Akreditif veya tahsil talimatı istiyorsa"),
    ("Banka tahsil talimatı", "odeme=vesaik mukabili", "Zorunlu", "İhracatçı / banka", "Belgelerin ihracatçının bankası aracılığıyla alıcının bankasına gönderilmesi"),
    ("Bitki Sağlık Sertifikası", "bitkisel urun=evet", "Zorunlu", "Tarım ve Orman il / ilçe müdürlüğü", "Bitki ve bitkisel ürünlerde"),
    ("Veteriner Sağlık Sertifikası", "hayvansal urun=evet", "Zorunlu", "Tarım ve Orman il / ilçe müdürlüğü", "Hayvan ve hayvansal ürünlerde"),
    ("ISPM 15 ısıl işlem işareti (ahşap ambalaj)", "ahsap ambalaj=evet", "Zorunlu", "Ambalaj üreticisi", "Ahşap palet / sandıkta ISPM 15 damgası"),
    ("Güvenlik Bilgi Formu (SDS) ve tehlikeli madde beyanı", "tehlikeli madde=evet", "Zorunlu", "İhracatçı / üretici",
     "Taşıma türüne göre IMDG (deniz), ADR (karayolu), IATA DGR (havayolu)"),
    ("Dahilde İşleme İzin Belgesi taahhüt kaydı", "dahilde isleme=evet", "Zorunlu", "İhracatçı / gümrük müşaviri", "Beyannamede DİİB numarası ve satır kodu"),
    ("Analiz / Sağlık Sertifikası", "gida=evet", "Önerilir", "Yetkili laboratuvar / kurum", "Alıcı ülke mevzuatı veya alıcı istiyorsa"),
]

TASIMA = {"deniz": "deniz", "denizyolu": "deniz", "deniz yolu": "deniz", "gemi": "deniz", "kara": "karayolu", "karayolu": "karayolu", "tir": "karayolu",
          "hava": "havayolu", "havayolu": "havayolu", "ucak": "havayolu", "demiryolu": "demiryolu", "tren": "demiryolu"}
ODEME = {"pesin": "pesin", "pesin odeme": "pesin", "mal mukabili": "mal mukabili", "acik hesap": "mal mukabili", "vesaik mukabili": "vesaik mukabili",
         "akreditif": "akreditif", "akreditifli": "akreditif", "kabul kredili": "kabul kredili"}
ES_ALANLAR = {"alici": "Alıcı", "consignee": "Alıcı", "gonderilen": "Alıcı", "toplam tutar": "Toplam Tutar", "tutar": "Toplam Tutar", "fatura tutari": "Toplam Tutar",
              "doviz": "Döviz", "para birimi": "Döviz", "koli sayisi": "Koli Sayısı", "kap adedi": "Koli Sayısı", "kap sayisi": "Koli Sayısı",
              "net agirlik": "Net Ağırlık", "brut agirlik": "Brüt Ağırlık", "gtip": "GTİP", "mense": "Menşe", "teslim sekli": "Teslim Şekli",
              "fatura no": "Fatura No", "fatura tarihi": "Fatura Tarihi", "yukleme limani": "Yükleme Limanı", "varis limani": "Varış Limanı",
              "mal tanimi": "Mal Tanımı", "gonderici": "Gönderici", "ihracatci": "Gönderici"}
ALAN_ADI = {"ulke grubu": "ülke grubu", "teslim": "teslim şekli", "tasima": "taşıma", "odeme": "ödeme", "urun turu": "ürün türü", "ahsap ambalaj": "ahşap ambalaj",
            "tehlikeli madde": "tehlikeli madde", "bitkisel urun": "bitkisel ürün", "hayvansal urun": "hayvansal ürün", "dahilde isleme": "dahilde işleme", "gida": "gıda"}
SAYISAL = {"Toplam Tutar", "Koli Sayısı", "Net Ağırlık", "Brüt Ağırlık"}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def sayi(x) -> Decimal | None:
    s = re.sub(r"[^\d.,-]", "", metin(x))
    if not s:
        return None
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


# ----------------------------------------------------------------------------
# Sevkiyat ve kurallar
# ----------------------------------------------------------------------------

def sevkiyat_oku(yol: Path) -> tuple[dict, list[dict]]:
    ham, uy = {}, []
    for r in tablo_oku(yol):
        if len(r) >= 2 and metin(r[0]) and katla(r[0]) != "alan":
            ham[metin(r[0])] = metin(r[1])
    k = {katla(a): v for a, v in ham.items()}
    s = {"ham": ham, "no": k.get("sevkiyat no", ""), "alici": k.get("alici", ""), "ulke": k.get("ulke", "")}
    grup = katla(k.get("ulke grubu", ""))
    if grup in ("ab", "avrupa birligi", "eu"):
        grup = "ab"
    elif grup in ("sta", "serbest ticaret anlasmasi"):
        grup = "sta"
    elif grup:
        grup = "diger"
    elif katla(s["ulke"]) in AB_ULKELERI:
        grup = "ab"
    else:
        grup = "diger"
        uy.append({"onem": "Orta", "tur": "Ülke grubu varsayıldı", "kim": s["ulke"] or "Ülke",
                   "aciklama": "Ülke AB üyesi değil ve 'Ülke Grubu' yazılmamış; 'Diğer' kabul edildi. Türkiye'nin serbest ticaret anlaşması olan bir ülkeyse "
                               "sevkiyat bilgisine 'Ülke Grubu;STA' yazın"})
    teslim = (k.get("teslim sekli", "").upper().split() or [""])[0]
    tasima = TASIMA.get(katla(k.get("tasima sekli", "")), katla(k.get("tasima sekli", "")))
    odeme = ODEME.get(katla(k.get("odeme sekli", "")), katla(k.get("odeme sekli", "")))
    urun = katla(k.get("urun turu", "")) or "sanayi"
    urun = {"sanayi urunu": "sanayi", "tarim urunu": "tarim", "islenmis tarim urunu": "islenmis tarim", "komur celik urunu": "komur celik", "akct": "komur celik"}.get(urun, urun)
    s["kosullar"] = {"ulke grubu": grup, "teslim": teslim.lower(), "tasima": tasima, "odeme": odeme, "urun turu": urun}
    for alan in ("bitkisel urun", "hayvansal urun", "ahsap ambalaj", "tehlikeli madde", "dahilde isleme", "gida"):
        s["kosullar"][alan] = "evet" if katla(k.get(alan, "")) in ("evet", "e", "var", "1") else "hayir"
    for a, v in k.items():                      # kullanıcı tanımlı ek alanlar (kendi kurallarında kullanılabilir)
        s["kosullar"].setdefault(a, katla(v))
    for alan, ad in (("teslim", "Teslim Şekli"), ("tasima", "Taşıma Şekli"), ("odeme", "Ödeme Şekli")):
        if not s["kosullar"][alan]:
            uy.append({"onem": "Yüksek", "tur": "Sevkiyat bilgisi eksik", "kim": ad, "aciklama": f"{ad} yazılmamış; buna bağlı evraklar listelenemedi"})
    if s["kosullar"]["teslim"] in ("fob", "fas", "cfr", "cif") and tasima and tasima != "deniz":
        uy.append({"onem": "Orta", "tur": "Teslim şekli / taşıma", "kim": teslim, "aciklama": f"{teslim} yalnız deniz taşımasında kullanılır; taşıma '{tasima}'"})
    return s, uy


def kural_oku(yol: Path | None) -> list[tuple]:
    kurallar = list(VARSAYILAN_KURALLAR)
    if not yol:
        return kurallar
    satirlar = tablo_oku(yol)
    b = [katla(c) for c in satirlar[0]]
    def al(r, *adlar):
        for a in adlar:
            if a in b and b.index(a) < len(r):
                return metin(r[b.index(a)])
        return ""
    for r in satirlar[1:]:
        evrak = al(r, "evrak", "belge")
        if not evrak:
            continue
        kosul = al(r, "kosul")
        kurallar = [k for k in kurallar if katla(k[0]) != katla(evrak)]          # aynı adlı varsayılan kural değiştirilir
        if katla(kosul) != "kaldir":
            kurallar.append((evrak, kosul or "her zaman", al(r, "zorunluluk") or "Zorunlu", al(r, "duzenleyen", "kim duzenler"), al(r, "aciklama")))
    return kurallar


def kosul_uyar(kosul: str, degerler: dict) -> tuple[bool, str]:
    k = kosul.strip().lower().translate(_TR)
    if katla(k) in ("her zaman", ""):
        return True, "Her ihracatta"
    nedenler = []
    for parca in k.split("&"):
        alan, _, deger = parca.partition("=")
        alan, secenekler = katla(alan), [katla(d) for d in deger.split("|")]
        mevcut = degerler.get(alan, "")
        if mevcut not in secenekler:
            return False, ""
        nedenler.append(f"{ALAN_ADI.get(alan, alan)}: {mevcut.upper() if alan in ('teslim', 'ulke grubu') else mevcut}")
    return True, "; ".join(nedenler)


# ----------------------------------------------------------------------------
# Evraklar ve tutarlılık
# ----------------------------------------------------------------------------

def evraklari_oku(yol: Path | None) -> dict[str, dict[str, str]]:
    evraklar = defaultdict(dict)
    if not yol:
        return evraklar
    for r in tablo_oku(yol):
        if len(r) >= 2 and metin(r[0]) and katla(r[0]) != "evrak":
            alan = metin(r[1]) if len(r) > 1 else ""
            evraklar[metin(r[0])][ES_ALANLAR.get(katla(alan), alan)] = metin(r[2]) if len(r) > 2 else ""
    return evraklar


def evrak_esle(gereken: str, mevcutlar: list[str]) -> str | None:
    """Gereken evrak adını hazırlanan evrak adlarıyla eşler (ortak anlamlı kelime)."""
    g = {w for w in katla(gereken).split() if len(w) > 2 and w not in {"veya", "belgesi", "sertifikasi", "senedi", "listesi"}}
    en, puan = None, 0
    for m in mevcutlar:
        ortak = len(g & set(katla(m).split()))
        if ortak > puan:
            en, puan = m, ortak
    return en


def ayni_mi(alan: str, a: str, b: str) -> bool:
    if alan in SAYISAL:
        x, y = sayi(a), sayi(b)
        return x is not None and y is not None and abs(x - y) <= Decimal("0.01")
    if alan == "GTİP":
        x, y = re.sub(r"\D", "", a), re.sub(r"\D", "", b)
        return x.startswith(y) or y.startswith(x)          # 6 haneli HS ile 12 haneli GTİP uyumlu sayılır
    if alan == "Teslim Şekli":
        return katla(a).split()[:1] == katla(b).split()[:1]     # "CIF Hamburg" ile "CIF" uyumlu
    return katla(a) == katla(b)


def kontrol_et(sevkiyat: dict, kurallar: list[tuple], evraklar: dict) -> dict:
    uy, liste = [], []
    mevcutlar = list(evraklar)
    kullanilan = set()
    for evrak, kosul, zorunluluk, duzenleyen, aciklama in kurallar:
        uyar, neden = kosul_uyar(kosul, sevkiyat["kosullar"])
        if not uyar:
            continue
        if any(katla(x["evrak"]) == katla(evrak) for x in liste):
            continue
        es = evrak_esle(evrak, mevcutlar) if evraklar else None
        if es:
            kullanilan.add(es)
        durum = "—" if not evraklar else ("Hazır" if es else "Eksik")
        liste.append({"evrak": evrak, "zorunluluk": zorunluluk, "neden": neden, "duzenleyen": duzenleyen, "aciklama": aciklama, "durum": durum, "eslesen": es or ""})
        if durum == "Eksik":
            uy.append({"onem": "Yüksek" if katla(zorunluluk) == "zorunlu" else "Bilgi", "tur": "Eksik evrak", "kim": evrak,
                       "aciklama": f"{zorunluluk}: {aciklama} ({neden}). Düzenleyen: {duzenleyen}"})
    for m in mevcutlar:
        if m not in kullanilan:
            uy.append({"onem": "Bilgi", "tur": "Listede olmayan evrak", "kim": m, "aciklama": "Hazırlanmış ama kural tablosuna göre bu sevkiyatta aranmıyor; alıcı talebiyse sorun değil"})

    alanlar = defaultdict(dict)
    for e, d in evraklar.items():
        for alan, deger in d.items():
            if deger:
                alanlar[alan][e] = deger
    tutarlilik = []
    for alan, d in alanlar.items():
        if len(d) < 2:
            tutarlilik.append({"alan": alan, "degerler": d, "durum": "Tek evrakta"})
            continue
        ref_evrak, ref = next(iter(d.items()))
        farkli = [e for e, v in d.items() if not ayni_mi(alan, ref, v)]
        tutarlilik.append({"alan": alan, "degerler": d, "durum": "Tutarsız" if farkli else "Tutarlı"})
        if farkli:
            uy.append({"onem": "Yüksek", "tur": "Tutarsızlık", "kim": alan,
                       "aciklama": "; ".join(f"{e}: {v}" for e, v in d.items()) + ". Evraklar arasında aynı olmalı (banka ve gümrük rezerv nedeni)"})
    # Sevkiyat bilgisiyle karşılaştırma
    for alan, anahtar in (("Alıcı", "alici"), ("Teslim Şekli", "teslim")):
        beklenen = sevkiyat["alici"] if anahtar == "alici" else sevkiyat["kosullar"]["teslim"]
        for e, v in alanlar.get(alan, {}).items():
            uygun = katla(v) == katla(beklenen) if anahtar == "alici" else katla(v).startswith(katla(beklenen))
            if beklenen and not uygun and not any(u["tur"] == "Tutarsızlık" and u["kim"] == alan for u in uy):
                uy.append({"onem": "Orta", "tur": "Sevkiyat bilgisinden farklı", "kim": alan, "aciklama": f"{e}: '{v}', sevkiyat bilgisi: '{beklenen}'"})
    net, brut = alanlar.get("Net Ağırlık", {}), alanlar.get("Brüt Ağırlık", {})
    for e in set(net) & set(brut):
        n, b = sayi(net[e]), sayi(brut[e])
        if n is not None and b is not None and n > b:
            uy.append({"onem": "Yüksek", "tur": "Ağırlık hatası", "kim": e, "aciklama": f"Net ağırlık ({net[e]}) brüt ağırlıktan ({brut[e]}) büyük"})
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    uy.sort(key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))
    return {"sevkiyat": sevkiyat, "liste": liste, "tutarlilik": tutarlilik, "evraklar": evraklar, "uyarilar": uy}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
UST = Alignment(vertical="top", wrap_text=True)
GRUP_ADI = {"ab": "AB", "sta": "STA ülkesi", "diger": "Diğer"}


def _baslik(ws, basliklar, genislik=()):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def rapor_yaz(cikti: Path, s: dict) -> None:
    sv, k = s["sevkiyat"], s["sevkiyat"]["kosullar"]
    wb = Workbook()
    kl = wb.active
    kl.title = "Kontrol Listesi"
    kl.append([f"İhracat evrak kontrol listesi — {sv['no']} · {sv['alici']} · {sv['ulke']} ({GRUP_ADI[k['ulke grubu']]})"])
    kl["A1"].font = Font(bold=True, size=12)
    kl.append([f"Teslim: {k['teslim'].upper() or '—'} · Taşıma: {k['tasima'] or '—'} · Ödeme: {k['odeme'] or '—'} · Ürün türü: {k['urun turu']}"])
    kl.append([])
    _baslik(kl, ["Evrak", "Zorunluluk", "Neden Gerekli", "Düzenleyen", "Açıklama", "Durum", "Eşleşen Evrak", "Nüsha / Not"], (46, 11, 30, 30, 60, 9, 26, 22))
    kl.freeze_panes = "A5"
    for x in s["liste"]:
        kl.append([x["evrak"], x["zorunluluk"], x["neden"], x["duzenleyen"], x["aciklama"], x["durum"], x["eslesen"], ""])
        kl.cell(kl.max_row, 5).alignment = kl.cell(kl.max_row, 3).alignment = UST
        renk = {"Hazır": "E3F4E1", "Eksik": "FDE2E1" if katla(x["zorunluluk"]) == "zorunlu" else "FFF4CE"}.get(x["durum"])
        if renk:
            kl.cell(kl.max_row, 6).fill = PatternFill("solid", fgColor=renk)
        kl.cell(kl.max_row, 8).fill = PatternFill("solid", fgColor="FFF4CE")

    tt = wb.create_sheet("Tutarlılık")
    adlar = list(s["evraklar"])
    _baslik(tt, ["Alan"] + adlar + ["Durum"], [18] + [26] * len(adlar) + [12])
    tt.freeze_panes = "B2"
    for x in sorted(s["tutarlilik"], key=lambda x: (x["durum"] != "Tutarsız", x["alan"])):
        tt.append([x["alan"]] + [x["degerler"].get(a, "") for a in adlar] + [x["durum"]])
        tt.cell(tt.max_row, len(adlar) + 2).fill = PatternFill("solid", fgColor={"Tutarsız": "FDE2E1", "Tutarlı": "E3F4E1"}.get(x["durum"], "FFFFFF"))

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Evrak / Alan", "Açıklama"], (9, 26, 40, 110))
    uy.freeze_panes = "A2"
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["kim"], u["aciklama"]])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(sevkiyat_yolu: Path, cikti: Path, evrak_yolu: Path | None = None, kural_yolu: Path | None = None) -> dict:
    sevkiyat, uy = sevkiyat_oku(sevkiyat_yolu)
    s = kontrol_et(sevkiyat, kural_oku(kural_yolu), evraklari_oku(evrak_yolu))
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    s["uyarilar"] = sorted(uy + s["uyarilar"], key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))
    rapor_yaz(cikti, s)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Sevkiyata göre gereken ihracat evraklarını listeler, evraklar arası tutarlılığı kontrol eder.")
    p.add_argument("--sevkiyat", type=Path, default=ORNEK / "sevkiyat.csv",
                   help="Alan;Değer: Sevkiyat No, Alıcı, Ülke, Ülke Grubu (AB/STA/Diğer), Teslim Şekli, Taşıma Şekli, Ödeme Şekli, Ürün Türü, "
                        "Bitkisel Ürün, Hayvansal Ürün, Ahşap Ambalaj, Tehlikeli Madde, Dahilde İşleme, Gıda")
    p.add_argument("--evraklar", type=Path, help="Hazırlanan evrakların bilgileri: Evrak; Alan; Değer (her alan bir satır)")
    p.add_argument("--kurallar", type=Path, help="Ek / değiştirilen kurallar: Evrak; Koşul; Zorunluluk; Düzenleyen; Açıklama (Koşul 'kaldır' → kural çıkarılır)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "ihracat_evrak_kontrol.xlsx")
    a = p.parse_args(argv)
    ornek = a.sevkiyat == ORNEK / "sevkiyat.csv"
    evrak = a.evraklar or (ORNEK / "evraklar.csv" if ornek else None)
    for y in (a.sevkiyat, evrak, a.kurallar):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        s = calistir(a.sevkiyat, a.cikti, evrak, a.kurallar)
    except (ValueError, IndexError) as h:
        print(f"[X] {h}")
        return 1
    k = s["sevkiyat"]["kosullar"]
    print(f"[OK] {s['sevkiyat']['ulke']} ({GRUP_ADI[k['ulke grubu']]}) · {k['teslim'].upper()} · {k['tasima']} · {k['odeme']}: {len(s['liste'])} evrak "
          f"({sum(katla(x['zorunluluk']) == 'zorunlu' for x in s['liste'])} zorunlu)")
    for x in s["liste"]:
        print(f"     [{x['durum']:<5}] {x['evrak']} ({x['zorunluluk']})")
    for u in s["uyarilar"]:
        if u["onem"] != "Bilgi":
            print(f"[!] {u['kim']} · {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
