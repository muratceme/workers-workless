import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def sevkiyat(tmp, **alanlar):
    yol = Path(tmp) / "s.csv"
    yol.write_text("Alan;Değer\n" + "\n".join(f"{k};{v}" for k, v in alanlar.items()) + "\n", encoding="utf-8")
    return yol


def evrak_adlari(s):
    return {x["evrak"].split(" (")[0] for x in s["liste"]}


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "i.xlsx"
        cls.s = main.calistir(ORNEK / "sevkiyat.csv", cls.cikti, ORNEK / "evraklar.csv")
        cls.u = {(u["tur"], u["kim"]) for u in cls.s["uyarilar"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_gereken_evraklar(self):
        ad = evrak_adlari(self.s)
        self.assertTrue({"Ticari Fatura", "Çeki Listesi", "Konşimento", "A.TR Dolaşım Belgesi", "Sigorta Poliçesi / Sertifikası", "Banka tahsil talimatı",
                         "ISPM 15 ısıl işlem işareti"} <= ad)
        self.assertFalse({"CMR Taşıma Senedi", "Menşe Şahadetnamesi", "EUR.1 / EUR-MED Dolaşım Sertifikası veya fatura beyanı",
                          "Akreditifte istenen belgeler"} & ad)
        d = {x["evrak"].split(" (")[0]: x["durum"] for x in self.s["liste"]}
        self.assertEqual((d["Ticari Fatura"], d["Konşimento"], d["Sigorta Poliçesi / Sertifikası"]), ("Hazır", "Hazır", "Eksik"))

    def test_tutarlilik(self):
        t = {x["alan"]: x["durum"] for x in self.s["tutarlilik"]}
        self.assertEqual((t["Alıcı"], t["Koli Sayısı"]), ("Tutarsız", "Tutarsız"))
        self.assertEqual((t["Brüt Ağırlık"], t["Toplam Tutar"], t["GTİP"], t["Teslim Şekli"]), ("Tutarlı", "Tutarlı", "Tutarlı", "Tutarlı"))
        self.assertTrue({("Tutarsızlık", "Alıcı"), ("Eksik evrak", "Sigorta Poliçesi / Sertifikası"), ("Listede olmayan evrak", "Kalite Sertifikası")} <= self.u)
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Kontrol Listesi", "Tutarlılık", "Uyarılar"])


class KuralTesti(unittest.TestCase):
    def test_ulke_ve_tasima(self):
        with tempfile.TemporaryDirectory() as tmp:
            c = Path(tmp) / "x.xlsx"
            s = main.calistir(sevkiyat(tmp, **{"Ülke": "Fransa", "Teslim Şekli": "FCA", "Taşıma Şekli": "Karayolu", "Ödeme Şekli": "Peşin", "Ürün Türü": "Tarım",
                                                 "Bitkisel Ürün": "Evet"}), c)
            ad = evrak_adlari(s)
            self.assertTrue({"CMR Taşıma Senedi", "EUR.1 / EUR-MED Dolaşım Sertifikası veya fatura beyanı", "Bitki Sağlık Sertifikası"} <= ad)
            self.assertFalse({"A.TR Dolaşım Belgesi", "Sigorta Poliçesi / Sertifikası", "Konşimento"} & ad)
            s = main.calistir(sevkiyat(tmp, **{"Ülke": "Birleşik Krallık", "Ülke Grubu": "STA", "Teslim Şekli": "CIP", "Taşıma Şekli": "Havayolu",
                                                 "Ödeme Şekli": "Akreditif", "Tehlikeli Madde": "Evet"}), c)
            ad = evrak_adlari(s)
            self.assertTrue({"Hava Yük Senedi", "EUR.1 / EUR-MED Dolaşım Sertifikası veya fatura beyanı", "Sigorta Poliçesi / Sertifikası",
                             "Akreditifte istenen belgeler", "Güvenlik Bilgi Formu"} <= ad)
            s = main.calistir(sevkiyat(tmp, **{"Ülke": "Örnek Ülke", "Teslim Şekli": "FOB", "Taşıma Şekli": "Karayolu", "Ödeme Şekli": "Mal mukabili"}), c)
            self.assertIn("Menşe Şahadetnamesi", evrak_adlari(s))
            turler = {u["tur"] for u in s["uyarilar"]}
            self.assertTrue({"Ülke grubu varsayıldı", "Teslim şekli / taşıma"} <= turler)

    def test_ek_kural_ve_karsilastirma(self):
        with tempfile.TemporaryDirectory() as tmp:
            k = Path(tmp) / "k.csv"
            k.write_text("Evrak;Koşul;Zorunluluk;Düzenleyen;Açıklama\nHelal Sertifikası;ülke=örnek ülke;Zorunlu;Yetkili kuruluş;Alıcı ülke şartı\n"
                         "İhracatçı Birliği kaydı / onayı;kaldır;;;\n", encoding="utf-8")
            s = main.calistir(sevkiyat(tmp, **{"Ülke": "Örnek Ülke", "Ülke Grubu": "Diğer", "Teslim Şekli": "EXW", "Taşıma Şekli": "Karayolu", "Ödeme Şekli": "Peşin"}),
                              Path(tmp) / "x.xlsx", kural_yolu=k)
            ad = evrak_adlari(s)
            self.assertIn("Helal Sertifikası", ad)
            self.assertNotIn("İhracatçı Birliği kaydı / onayı", ad)
        self.assertTrue(main.ayni_mi("Toplam Tutar", "85.512,00", "85512"))
        self.assertTrue(main.ayni_mi("GTİP", "6302.60", "630260000011"))
        self.assertFalse(main.ayni_mi("Koli Sayısı", "431", "430"))
        self.assertTrue(main.ayni_mi("Alıcı", "EXAMPLE gmbh", "Example GmbH"))

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--sevkiyat", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
