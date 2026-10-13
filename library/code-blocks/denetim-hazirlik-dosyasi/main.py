"""
Denetim Hazırlık Dosyası — Workers / Workless kod bloğu
Muhasebe › Muhasebe Müdürü

Bağımsız denetçinin talep listesindeki belge ve mutabakatların hazırlık durumunu izler.
  Kapsam             = tüm talepler − "Uygulanmaz" olanlar
  Hazırlık oranı     = (Gönderildi + Hazır) ÷ kapsam
  Gecikti            = gönderilmemiş ve istenen tarihi --tarih gününden önce;  gecikme = takvim günü farkı
  Yaklaşıyor         = gönderilmemiş, hazır değil ve istenen tarihe --gun (3) gün veya daha az kalmış
  Gecikmeli teslim   = gönderim tarihi istenen tarihten sonra
  Dosya kontrolü     = --klasor verilirse "Dosya" sütunundaki adlar klasörde (alt klasörler dahil) aranır
Kod belgelerin içeriğini okumaz ve yeterliliğini değerlendirmez; yalnız durumu ve dosyanın varlığını izler.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
SUTUN = {"no": ("no", "talep no", "sira", "sira no"), "talep": ("talep", "belge", "istenen belge", "talep edilen belge", "aciklama"), "alan": ("alan", "hesap alani", "bolum", "kategori"), "sorumlu": ("sorumlu", "sorumlu kisi", "hazirlayan"),
         "termin": ("istenen tarih", "termin", "son tarih", "teslim tarihi"), "durum": ("durum",), "gonderim": ("gonderim tarihi", "gonderildi", "teslim edildigi tarih"), "dosya": ("dosya", "dosya adi", "referans"),
         "not": ("not", "notlar")}
DURUM = {"gonderildi": "Gönderildi", "teslim edildi": "Gönderildi", "iletildi": "Gönderildi", "hazir": "Hazır", "tamamlandi": "Hazır", "hazirlaniyor": "Hazırlanıyor", "devam": "Hazırlanıyor", "devam ediyor": "Hazırlanıyor",
         "": "Başlamadı", "baslamadi": "Başlamadı", "bekliyor": "Başlamadı", "uygulanmaz": "Uygulanmaz", "yok": "Uygulanmaz", "gecerli degil": "Uygulanmaz"}
IZ_RENK = {"Gecikti": KIRMIZI, "Yaklaşıyor": SARI, "Hazır": MAVI, "Teslim edildi": YESIL, "Gecikmeli teslim": YESIL}


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("talep",)):
        ad = metin(r["talep"])
        if not ad:
            continue
        d = katla(r["durum"])
        if d not in DURUM:
            uy.append(uyari("Orta", "Durum tanınmadı", ad, f"'{metin(r['durum'])}' durumu tanınmadı; talep 'Başlamadı' sayıldı. Gönderildi, Hazır, Hazırlanıyor, Başlamadı veya Uygulanmaz yazın."))
        if metin(r["termin"]) and tarih(r["termin"]) is None:
            uy.append(uyari("Orta", "Tarih okunamadı", ad, f"İstenen tarih '{metin(r['termin'])}' okunamadı; talep terminsiz izlenir."))
        no = metin(r["no"]) or str(len(l) + 1)
        l.append({"no": no[:-2] if no.endswith(".0") else no, "talep": ad, "alan": metin(r["alan"]) or "(alan yok)", "sorumlu": metin(r["sorumlu"]), "termin": tarih(r["termin"]), "durum": DURUM.get(d, "Başlamadı"),
                  "gonderim": tarih(r["gonderim"]), "dosyalar": [x.strip() for x in re.split(r"[,|]", metin(r["dosya"])) if x.strip()], "not": metin(r["not"])})
    return l, uy


def _ozet(talepler: list[dict], alan: str) -> list[dict]:
    o = {}
    for t in talepler:
        if t["durum"] == "Uygulanmaz":
            continue
        x = o.setdefault(katla(t[alan]), {"ad": t[alan] or "(sorumlu yok)", "toplam": 0, "gonderildi": 0, "hazir": 0, "acik": 0, "geciken": 0})
        x["toplam"] += 1
        x["gonderildi" if t["durum"] == "Gönderildi" else "hazir" if t["durum"] == "Hazır" else "acik"] += 1
        x["geciken"] += t["izleme"] == "Gecikti"
    return sorted(({**x, "oran": yuvarla(Decimal(100 * (x["gonderildi"] + x["hazir"])) / x["toplam"])} for x in o.values()), key=lambda x: (-x["geciken"], x["oran"], katla(x["ad"])))


def degerlendir(talepler: list[dict], bugun: date, gun: int = 3, klasor: Path | None = None) -> dict:
    if not talepler:
        raise ValueError("Listede izlenecek denetim talebi yok.")
    uy, gorulen = [], set()
    mevcut = {p.name.casefold(): p.name for p in klasor.rglob("*") if p.is_file()} if klasor else {}
    kullanilan = set()
    for t in talepler:
        if t["no"] in gorulen:
            uy.append(uyari("Orta", "Talep numarası tekrar ediyor", t["no"], "Aynı numara birden çok talepte var."))
        gorulen.add(t["no"])
        t["gecikme"], t["kalan"] = 0, (t["termin"] - bugun).days if t["termin"] else None
        if t["durum"] == "Uygulanmaz":
            t["izleme"] = "Uygulanmaz"
        elif t["durum"] == "Gönderildi":
            t["izleme"] = "Gecikmeli teslim" if t["termin"] and t["gonderim"] and t["gonderim"] > t["termin"] else "Teslim edildi"
            if t["izleme"] == "Gecikmeli teslim":
                uy.append(uyari("Bilgi", "Gecikmeli teslim", t["talep"], f"İstenen {t['termin']:%d.%m.%Y}, gönderim {t['gonderim']:%d.%m.%Y}."))
        elif t["termin"] and t["termin"] < bugun:
            t["izleme"], t["gecikme"] = "Gecikti", -t["kalan"]
            uy.append(uyari("Yüksek", "Geciken talep", t["talep"], f"İstenen tarih {t['termin']:%d.%m.%Y}, {t['gecikme']} gün gecikmede. Sorumlu: {t['sorumlu'] or 'belirtilmemiş'}."
                            + (" Belge hazır görünüyor; denetçiye gönderin." if t["durum"] == "Hazır" else "")))
        elif t["durum"] == "Hazır":
            t["izleme"] = "Hazır"
        elif not t["termin"]:
            t["izleme"] = "Termin yok"
        else:
            t["izleme"] = "Yaklaşıyor" if t["kalan"] <= gun else "Bekliyor"
        if not t["sorumlu"] and t["durum"] not in ("Uygulanmaz", "Gönderildi"):
            uy.append(uyari("Orta", "Sorumlu yok", t["talep"], "Talebe sorumlu atanmamış."))
        t["eksik_dosya"] = []
        if klasor and t["durum"] != "Uygulanmaz":
            t["eksik_dosya"] = [d for d in t["dosyalar"] if d.casefold() not in mevcut]
            kullanilan |= {d.casefold() for d in t["dosyalar"]}
            if t["durum"] in ("Gönderildi", "Hazır"):
                if t["eksik_dosya"]:
                    uy.append(uyari("Yüksek", "Dosya klasörde yok", t["talep"], f"Durum '{t['durum']}' ama şu dosya(lar) klasörde bulunamadı: {', '.join(t['eksik_dosya'])}."))
                elif not t["dosyalar"]:
                    uy.append(uyari("Orta", "Dosya adı yazılmamış", t["talep"], f"Durum '{t['durum']}' ama Dosya sütunu boş; belgenin klasörde olduğu doğrulanamadı."))
            elif t["dosyalar"] and not t["eksik_dosya"]:
                uy.append(uyari("Bilgi", "Dosya var, durum açık", t["talep"], f"Yazılan dosyaların tümü klasörde ama durum '{t['durum']}'. Belge tamamsa durumu güncelleyin."))
    if klasor:
        fazla = sorted(ad for k, ad in mevcut.items() if k not in kullanilan)
        if fazla:
            uy.append(uyari("Bilgi", "Talebe bağlanmamış dosya", f"{len(fazla)} dosya", "Klasörde olup hiçbir talepte adı geçmeyen dosyalar: " + ", ".join(fazla[:15]) + (" …" if len(fazla) > 15 else "")))
    kapsam = [t for t in talepler if t["durum"] != "Uygulanmaz"]
    say = {d: sum(t["durum"] == d for t in kapsam) for d in ("Gönderildi", "Hazır", "Hazırlanıyor", "Başlamadı")}
    return {"talepler": talepler, "kapsam": len(kapsam), "say": say, "oran": yuvarla(Decimal(100 * (say["Gönderildi"] + say["Hazır"])) / len(kapsam)) if kapsam else None,
            "geciken": sum(t["izleme"] == "Gecikti" for t in talepler), "yaklasan": sum(t["izleme"] == "Yaklaşıyor" for t in talepler), "alanlar": _ozet(talepler, "alan"), "sorumlular": _ozet(talepler, "sorumlu"),
            "bugun": bugun, "klasor": klasor, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Talep Listesi"
    _baslik(ws, ["No", "Talep", "Alan", "Sorumlu", "İstenen Tarih", "Durum", "İzleme", "Gecikme (gün)", "Gönderim", "Dosya", "Klasörde Eksik", "Not"], (6, 46, 24, 18, 12, 13, 17, 9, 12, 28, 24, 30))
    for t in r["talepler"]:
        _satir(ws, [t["no"], t["talep"], t["alan"], t["sorumlu"], t["termin"], t["durum"], t["izleme"], t["gecikme"] or None, t["gonderim"], ", ".join(t["dosyalar"]), ", ".join(t["eksik_dosya"]), t["not"]], tarih_s=(5, 9),
               dolgu={7: IZ_RENK.get(t["izleme"]), 11: KIRMIZI if t["eksik_dosya"] and t["durum"] in ("Gönderildi", "Hazır") else None}, sar=(2,))
    ws.append([])
    s = r["say"]
    ws.append([f"Durum tarihi {r['bugun']:%d.%m.%Y} · kapsam {r['kapsam']} talep · gönderildi {s['Gönderildi']}, hazır {s['Hazır']}, hazırlanıyor {s['Hazırlanıyor']}, başlamadı {s['Başlamadı']}"
               + (f" · hazırlık oranı %{r['oran']}" if r["oran"] is not None else "") + ("" if r["klasor"] else " · dosya kontrolü yapılmadı (--klasor verilmedi)")])
    for ad, veri, baslik in (("Alan Özeti", r["alanlar"], "Alan"), ("Sorumlu Özeti", r["sorumlular"], "Sorumlu")):
        ws = wb.create_sheet(ad)
        _baslik(ws, [baslik, "Talep", "Gönderildi", "Hazır", "Açık", "Geciken", "Hazırlık %"], (32, 8, 11, 8, 8, 9, 12))
        for x in veri:
            _satir(ws, [x["ad"], x["toplam"], x["gonderildi"] or None, x["hazir"] or None, x["acik"] or None, x["geciken"] or None, x["oran"]], para_s=(7,), dolgu={6: KIRMIZI if x["geciken"] else None})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, gun: int = 3, klasor: Path | None = None) -> dict:
    talepler, uy = oku(yol)
    r = degerlendir(talepler, bugun, gun, klasor)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Denetim talep listesindeki belgelerin hazırlık durumunu izler.")
    p.add_argument("--talepler", type=Path, help="No, Talep, Alan, Sorumlu, İstenen Tarih, Durum, Gönderim Tarihi, Dosya, Not")
    p.add_argument("--klasor", type=Path, help="Denetim belgelerinin durduğu klasör; verilirse Dosya sütunundaki adlar burada aranır")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--gun", type=int, default=3, help="İstenen tarihe kaç gün kala 'Yaklaşıyor' sayılsın (varsayılan 3)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "denetim_hazirlik.xlsx")
    a = p.parse_args(argv)
    ornek = a.talepler is None
    yol = ORNEK / "talepler.csv" if ornek else a.talepler
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    if bugun is None or not 0 <= a.gun <= 90:
        print("[X] --tarih GG.AA.YYYY biçiminde, --gun 0–90 arasında olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    if a.klasor and not a.klasor.is_dir():
        print(f"[X] Klasör bulunamadı: {a.klasor}")
        return 1
    try:
        r = calistir(yol, a.cikti, bugun, a.gun, a.klasor)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    s = r["say"]
    print(f"[OK] {r['kapsam']} talep · gönderildi {s['Gönderildi']}, hazır {s['Hazır']}, hazırlanıyor {s['Hazırlanıyor']}, başlamadı {s['Başlamadı']} · geciken {r['geciken']}, yaklaşan {r['yaklasan']}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
