import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)


def t(no, durum="Başlamadı", termin=date(2026, 10, 20), gonderim=None, dosya=(), sorumlu="S", alan="A"):
    return {"no": str(no), "talep": f"Talep {no}", "alan": alan, "sorumlu": sorumlu, "termin": termin, "durum": durum, "gonderim": gonderim, "dosyalar": list(dosya), "not": ""}


class DegerlendirmeTesti(unittest.TestCase):
    def test_izleme(self):
        l = [t(1, "Gönderildi", date(2026, 10, 5), date(2026, 10, 5)), t(2, "Gönderildi", date(2026, 10, 5), date(2026, 10, 6)), t(3, "Hazır", date(2026, 10, 9)), t(4, "Hazır", date(2026, 10, 12)),
             t(5, "Hazırlanıyor", date(2026, 10, 11)), t(6, termin=date(2026, 10, 12)), t(7, termin=date(2026, 10, 15)), t(8, termin=date(2026, 10, 16)), t(9, termin=None, sorumlu=""), t(10, "Uygulanmaz", sorumlu=""),
             t(11, "Gönderildi", None, None, sorumlu="")]
        r = main.degerlendir(l, BUGUN)
        self.assertEqual([(x["izleme"], x["gecikme"]) for x in r["talepler"]],
                         [("Teslim edildi", 0), ("Gecikmeli teslim", 0), ("Gecikti", 3), ("Hazır", 0), ("Gecikti", 1), ("Yaklaşıyor", 0), ("Yaklaşıyor", 0), ("Bekliyor", 0), ("Termin yok", 0), ("Uygulanmaz", 0), ("Teslim edildi", 0)])
        self.assertEqual((r["kapsam"], r["say"], r["oran"], r["geciken"], r["yaklasan"]), (10, {"Gönderildi": 3, "Hazır": 2, "Hazırlanıyor": 1, "Başlamadı": 4}, D("50.00"), 2, 2))
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]), [("Geciken talep", "Talep 3"), ("Geciken talep", "Talep 5"), ("Gecikmeli teslim", "Talep 2"), ("Sorumlu yok", "Talep 9")])
        self.assertEqual(main.degerlendir([t(7, termin=date(2026, 10, 15))], BUGUN, 2)["talepler"][0]["izleme"], "Bekliyor")

    def test_ozet(self):
        l = [t(1, "Gönderildi", date(2026, 10, 5), date(2026, 10, 5), alan="Banka", sorumlu="Ali"), t(2, "Hazırlanıyor", date(2026, 10, 9), alan="Banka", sorumlu="Ali"), t(3, "Hazır", alan="Stok", sorumlu="Ece"),
             t(4, alan="Stok", sorumlu="Ece"), t(1, "Uygulanmaz", alan="Vergi", sorumlu="Ece"), t(6, "Gönderildi", alan="Vergi", sorumlu="ece")]
        r = main.degerlendir(l, BUGUN)
        self.assertEqual([(x["ad"], x["toplam"], x["gonderildi"], x["hazir"], x["acik"], x["geciken"], x["oran"]) for x in r["alanlar"]],
                         [("Banka", 2, 1, 0, 1, 1, D("50.00")), ("Stok", 2, 0, 1, 1, 0, D("50.00")), ("Vergi", 1, 1, 0, 0, 0, D("100.00"))])
        self.assertEqual([(x["ad"], x["toplam"], x["oran"]) for x in r["sorumlular"]], [("Ali", 2, D("50.00")), ("Ece", 3, D("66.67"))])
        self.assertIn("Talep numarası tekrar ediyor", [u["tur"] for u in r["uyarilar"]])
        self.assertEqual(main.degerlendir([t(1, "Uygulanmaz")], BUGUN)["oran"], None)
        with self.assertRaises(ValueError):
            main.degerlendir([], BUGUN)

    def test_klasor(self):
        with tempfile.TemporaryDirectory() as tmp:
            k = Path(tmp)
            (k / "alt").mkdir()
            for ad in ("mizan.xlsx", "alt/Banka.PDF", "fazla.txt", "stok.xlsx"):
                (k / ad).write_text("x", encoding="utf-8")
            l = [t(1, "Gönderildi", gonderim=BUGUN, dosya=["MIZAN.xlsx"]), t(2, "Hazır", dosya=["banka.pdf", "ekstre.pdf"]), t(3, "Hazır"), t(4, "Hazırlanıyor", dosya=["stok.xlsx"]), t(5, dosya=["yok.xlsx"]),
                 t(6, "Uygulanmaz", dosya=["fazla.txt"])]
            r = main.degerlendir(l, BUGUN, 3, k)
            self.assertEqual([x["eksik_dosya"] for x in r["talepler"]], [[], ["ekstre.pdf"], [], [], ["yok.xlsx"], []])
            self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]),
                             [("Dosya adı yazılmamış", "Talep 3"), ("Dosya klasörde yok", "Talep 2"), ("Dosya var, durum açık", "Talep 4"), ("Talebe bağlanmamış dosya", "1 dosya")])
            self.assertIn("fazla.txt", next(u["aciklama"] for u in r["uyarilar"] if u["tur"] == "Talebe bağlanmamış dosya"))


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "d.xlsx"
            r = main.calistir(ORNEK / "talepler.csv", cikti, BUGUN)
            self.assertEqual([(x["no"], x["izleme"], x["gecikme"]) for x in r["talepler"]],
                             [("1", "Teslim edildi", 0), ("2", "Gecikmeli teslim", 0), ("3", "Gecikti", 3), ("4", "Gecikti", 3), ("5", "Yaklaşıyor", 0), ("6", "Hazır", 0), ("7", "Yaklaşıyor", 0), ("8", "Yaklaşıyor", 0),
                              ("9", "Hazır", 0), ("10", "Bekliyor", 0), ("11", "Bekliyor", 0), ("12", "Bekliyor", 0), ("13", "Uygulanmaz", 0)])
            self.assertEqual((r["kapsam"], r["say"], r["oran"], r["geciken"], r["yaklasan"]), (12, {"Gönderildi": 2, "Hazır": 3, "Hazırlanıyor": 3, "Başlamadı": 4}, D("41.67"), 2, 3))
            self.assertEqual([(u["onem"], u["tur"]) for u in r["uyarilar"]], [("Yüksek", "Geciken talep"), ("Yüksek", "Geciken talep"), ("Orta", "Sorumlu yok"), ("Bilgi", "Gecikmeli teslim")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Talep Listesi", "Alan Özeti", "Sorumlu Özeti", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "t.csv"
            y.write_text("Belge;Bölüm;Hazırlayan;Termin;Durum;Dosya Adı\nMizan;Genel;Ali;05.10.2026;TESLİM EDİLDİ;a.xlsx, b.pdf | c.pdf\nStok;Stok;Ece;x;bitti;\n;;;;;\nKredi;;;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["talep"], x["alan"], x["termin"], x["durum"], x["dosyalar"]) for x in l], [u["tur"] for u in uy]),
                             ([("1", "Mizan", "Genel", date(2026, 10, 5), "Gönderildi", ["a.xlsx", "b.pdf", "c.pdf"]), ("2", "Stok", "Stok", None, "Başlamadı", []), ("3", "Kredi", "(alan yok)", None, "Başlamadı", [])],
                              ["Durum tanınmadı", "Tarih okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--talepler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "12.10.2026", "--gun", "5", "--klasor", tmp]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual((main.main(arg + ["--gun", "-1"]), main.main(arg + ["--tarih", "x"])), (2, 2))
            self.assertEqual((main.main(["--talepler", str(Path(tmp) / "yok.csv")]), main.main(arg + ["--klasor", str(Path(tmp) / "yok")])), (1, 1))
            for icerik in ("Sorumlu;Durum\nA;\n", "Talep;Durum\n;\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
