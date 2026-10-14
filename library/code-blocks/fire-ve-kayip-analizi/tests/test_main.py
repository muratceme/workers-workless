import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def f(grup, neden, tutar):
    return {"tarih": date(2026, 9, 1), "grup": grup, "neden": neden or main.YOK, "adet": None, "tutar": D(tutar)}


def gr(grup, satis, fark=None):
    return {"grup": grup, "satis": D(satis), "fark": None if fark is None else D(fark)}


def veri():
    return [f("G1", "SKT", 300), f("g1", "Kırık", 100), f("G2", "skt", 200), f("G2", "Hırsızlık", 50), f("G3", "", 40)], [gr("G1", 20000, -100), gr("G2", 5000, -300), gr("G4", 1000, 20)]


def k(s):
    return (s["grup"], s["satis"], s["fire"], s["fire_oran"], s["bilinmeyen"], s["toplam"], s["kayip_oran"], s["bilinmeyen_pay"], s["kayit"])


class AnalizTesti(unittest.TestCase):
    def test_kayip(self):
        r = main.analiz(*veri())
        self.assertEqual([k(s) for s in r["gruplar"]],
                         [("G2", D(5000), D(250), D("5.00"), D(300), D(550), D("11.00"), D("54.55"), 2), ("G1", D(20000), D(400), D("2.00"), D(100), D(500), D("2.50"), D("20.00"), 2), ("G4", D(1000), D(0), D("0.00"), D(-20), D(-20), D("-2.00"), None, 0),
                          ("G3", None, D(40), None, None, D(40), None, None, 1)])
        self.assertEqual(r["ozet"], {"kayit": 5, "satis": D(26000), "fire": D(690), "bilinmeyen": D(380), "toplam": D(1070), "fire_oran": D("2.65"), "kayip_oran": D("4.12")})
        self.assertEqual([(n["ad"], n["sinif"], n["kayit"], n["tutar"], n["pay"], n["kum"]) for n in r["nedenler"]],
                         [("SKT", "A", 2, D(500), D("72.46"), D("72.46")), ("Kırık", "A", 1, D(100), D("14.49"), D("86.96")), ("Hırsızlık", "B", 1, D(50), D("7.25"), D("94.20")), (main.YOK, "B", 1, D(40), D("5.80"), D("100.00"))])
        self.assertEqual([(m["grup"], m["hucre"], m["toplam"]) for m in r["matris"]], [("G2", [D(200), D(0), D(50), D(0)], D(250)), ("G1", [D(300), D(100), D(0), D(0)], D(400)), ("G3", [D(0), D(0), D(0), D(40)], D(40))])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Bilgi", "Grubun satışı yok", "G3"), ("Bilgi", "Kaybın çoğu kayıt dışı", "G2"), ("Orta", "Kayıp oranı yüksek", "G1"), ("Orta", "Nedeni yazılmamış fire", main.YOK), ("Yüksek", "Kayıp oranı yüksek", "G2")])

    def test_esik_ve_grupsuz(self):
        u = lambda e: sorted((x["onem"], x["kim"]) for x in main.analiz(*veri(), esik=D(e))["uyarilar"] if x["tur"] == "Kayıp oranı yüksek")   # noqa: E731
        self.assertEqual((u("2.5"), u("5.5"), u("11")), ([("Yüksek", "G2")], [("Orta", "G2")], []))
        r = main.analiz(veri()[0])
        self.assertEqual(([(s["grup"], s["fire"], s["fire_pay"], s["kayip_oran"]) for s in r["gruplar"]], r["ozet"], [x["tur"] for x in r["uyarilar"]]),
                         ([("G1", D(400), D("57.97"), None), ("G2", D(250), D("36.23"), None), ("G3", D(40), D("5.80"), None)],
                          {"kayit": 5, "satis": None, "fire": D(690), "bilinmeyen": None, "toplam": D(690), "fire_oran": None, "kayip_oran": None}, ["Nedeni yazılmamış fire"]))
        self.assertEqual(main.analiz([], [gr("G", 100, -5)])["ozet"]["kayip_oran"], D("5.00"))
        with self.assertRaises(ValueError):
            main.analiz([], [])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "fire_kayitlari.csv", cikti, ORNEK / "urun_gruplari.csv")
            self.assertEqual([k(s) for s in r["gruplar"]],
                             [("Meyve Sebze", D(480000), D(19200), D("4.00"), D(2400), D(21600), D("4.50"), D("11.11"), 2), ("Kozmetik", D(350000), D(2100), D("0.60"), D(9800), D(11900), D("3.40"), D("82.35"), 2),
                              ("Unlu Mamul", D(210000), D(6300), D("3.00"), D(300), D(6600), D("3.14"), D("4.55"), 1), ("Et Tavuk", D(620000), D(9300), D("1.50"), D(1800), D(11100), D("1.79"), D("16.22"), 2),
                              ("Süt Ürünleri", D(540000), D(6750), D("1.25"), D(900), D(7650), D("1.42"), D("11.76"), 2), ("Temizlik", D(400000), D(1200), D("0.30"), D(-600), D(600), D("0.15"), None, 1)])
            self.assertEqual(r["ozet"], {"kayit": 10, "satis": D(2600000), "fire": D(44850), "bilinmeyen": D(14600), "toplam": D(59450), "fire_oran": D("1.73"), "kayip_oran": D("2.29")})
            self.assertEqual([(n["ad"], n["tutar"], n["sinif"]) for n in r["nedenler"]],
                             [("Bozulma", D(14400), "A"), ("SKT", D(11600), "A"), ("Gün sonu imha", D(6300), "A"), ("Ezilme", D(4800), "A"), ("Ambalaj hasarı", D(3250), "B"), ("Soğuk zincir", D(3100), "B"), ("Hırsızlık (tespit edilen)", D(1400), "B")])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Yüksek", "Kayıp oranı yüksek", "Meyve Sebze"), ("Orta", "Kayıp oranı yüksek", "Kozmetik"), ("Orta", "Kayıp oranı yüksek", "Unlu Mamul"), ("Bilgi", "Kaybın çoğu kayıt dışı", "Kozmetik")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Ürün Grupları", "Neden Pareto", "Grup × Neden", "Fire Kayıtları", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            fy, gy = Path(tmp) / "f.csv", Path(tmp) / "g.csv"
            fy.write_text("Kayıt Tarihi;Reyon;Fire Nedeni;Miktar;Fire Tutarı\n01.09.2026;Süt;SKT;1.200;1.250,50\n;Süt;;;10\n;Süt;X;;y\n;Süt;X;;-1\n;;X;;5\n", encoding="utf-8")
            gy.write_text("Kategori;Ciro;Envanter Farkı\nSüt;100.000;-1.500,50\nEt;50;\nsüt;5;\nBal;x;\nÇay;10;y\n;10;\n", encoding="utf-8")
            fl, uy = main.oku_fire(fy)
            self.assertEqual(([(x["tarih"], x["grup"], x["neden"], x["adet"], x["tutar"]) for x in fl], [u["tur"] for u in uy]),
                             ([(date(2026, 9, 1), "Süt", "SKT", D(1200), D("1250.50")), (None, "Süt", main.YOK, None, D(10))], ["Fire kaydı okunamadı"] * 2))
            gl, uy = main.oku_gruplar(gy)
            self.assertEqual((gl, [(u["tur"], u["kim"]) for u in uy]),
                             ([{"grup": "Süt", "satis": D(100000), "fark": D("-1500.50")}, {"grup": "Et", "satis": D(50), "fark": None}], [("Grup satırı okunamadı", "süt"), ("Grup satırı okunamadı", "Bal"), ("Grup satırı okunamadı", "Çay")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--fire", str(fy), "--cikti", c]
            self.assertEqual((main.main(arg), main.main(arg + ["--gruplar", str(gy), "--esik", "1,5"])), (0, 0))
            self.assertEqual([main.main(arg + v) for v in (["--esik", "0"], ["--esik", "x"])], [2] * 2)
            self.assertEqual((main.main(["--fire", str(Path(tmp) / "yok.csv")]), main.main(arg + ["--gruplar", str(Path(tmp) / "yok.csv")])), (1, 1))
            for icerik in ("Ürün Grubu;Neden\nA;B\n", "Ürün Grubu;Tutar\n;5\n"):
                fy.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
