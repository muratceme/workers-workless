"""
Fire ve Kayıp Analizi — Workers / Workless kod bloğu
Mağaza Operasyonları › Mağaza Müdürü

Fire kayıtlarını ürün grubu ve neden bazında analiz eder; sayım farklarıyla birlikte kayıp oranını hesaplar.
  Kayıtlı fire        = fire kayıtlarındaki tutarların toplamı (bozulma, SKT, hasar, imha …)
  Bilinmeyen kayıp    = −sayım farkı      (sayım farkı eksi ise kayıp; artı ise sayım fazlası, kaybı azaltır)
  Toplam kayıp        = kayıtlı fire + bilinmeyen kayıp
  Fire % / kayıp %    = kayıtlı fire ya da toplam kayıp ÷ satış × 100
  Neden Pareto        = fire tutarına göre nedenlerin payı ve kümülatif payı;   A sınıfı: önceki kümülatif pay %80'in altında
  Uyarı               = grubun kayıp oranı --esik (%2) üstünde (iki katını aşarsa yüksek) · bilinmeyen kayıp toplamın yarısından fazla
Fire, sayım farkı ve satış aynı esasta (hepsi maliyet ya da hepsi satış bedeli) verilmelidir; kod birini diğerine çevirmez.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

F_SUTUN = {"tarih": ("tarih", "kayit tarihi"), "grup": ("urun grubu", "kategori", "reyon", "grup"), "neden": ("neden", "fire nedeni", "imha nedeni"), "adet": ("adet", "miktar"), "tutar": ("tutar", "fire tutari", "maliyet")}
G_SUTUN = {"grup": ("urun grubu", "kategori", "reyon", "grup"), "satis": ("satis", "satis tutari", "ciro"), "fark": ("sayim farki", "sayim farki tutari", "envanter farki")}
YOK = "(neden yazılmamış)"
SIFIR = Decimal(0)


def oku_fire(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, F_SUTUN, ("grup", "tutar")):
        grup = metin(r["grup"])
        if not grup:
            continue
        tutar = para(r["tutar"])
        if tutar is None or tutar < 0:
            uy.append(uyari("Yüksek", "Fire kaydı okunamadı", grup, f"{yol.name} satır {r['_satir']}: tutar '{metin(r['tutar'])}' okunamadı. Satır alınmadı."))
            continue
        l.append({"tarih": tarih(r["tarih"]), "grup": grup, "neden": metin(r["neden"]) or YOK, "adet": para(r["adet"]) if metin(r["adet"]) else None, "tutar": tutar})
    return l, uy


def oku_gruplar(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, G_SUTUN, ("grup", "satis")):
        grup = metin(r["grup"])
        if not grup:
            continue
        satis, fark = para(r["satis"]), para(r["fark"]) if metin(r["fark"]) else None
        if satis is None or satis < 0 or (metin(r["fark"]) and fark is None) or katla(grup) in gorulen:
            uy.append(uyari("Yüksek", "Grup satırı okunamadı", grup, f"{yol.name} satır {r['_satir']}: " + ("aynı ürün grubu daha önce geçti" if katla(grup) in gorulen else "satış ya da sayım farkı okunamadı") + ". Satır alınmadı."))
            continue
        gorulen.add(katla(grup))
        l.append({"grup": grup, "satis": satis, "fark": fark})
    return l, uy


def _oran(pay: Decimal, payda: Decimal | None) -> Decimal | None:
    return yuvarla(pay * 100 / payda) if payda else None


def analiz(fire: list[dict], gruplar: list[dict] | None = None, esik: Decimal = Decimal(2)) -> dict:
    gruplar = gruplar or []
    if not fire and not gruplar:
        raise ValueError("Analiz edilecek fire kaydı ya da grup verisi yok.")
    uy, g = [], {}
    for x in gruplar:
        g[katla(x["grup"])] = {"grup": x["grup"], "satis": x["satis"], "fark": x["fark"], "fire": SIFIR, "kayit": 0}
    for f in fire:
        s = g.setdefault(katla(f["grup"]), {"grup": f["grup"], "satis": None, "fark": None, "fire": SIFIR, "kayit": 0})
        s["fire"], s["kayit"] = s["fire"] + f["tutar"], s["kayit"] + 1
    top_fire = sum((s["fire"] for s in g.values()), SIFIR)
    for s in g.values():
        s["bilinmeyen"] = -s["fark"] if s["fark"] is not None else None
        s["toplam"] = s["fire"] + (s["bilinmeyen"] or SIFIR)
        s["fire_oran"], s["kayip_oran"], s["fire_pay"] = _oran(s["fire"], s["satis"]), _oran(s["toplam"], s["satis"]), _oran(s["fire"], top_fire)
        s["bilinmeyen_pay"] = _oran(s["bilinmeyen"], s["toplam"]) if s["bilinmeyen"] is not None and s["bilinmeyen"] > 0 and s["toplam"] > 0 else None
        if gruplar and s["satis"] is None:
            uy.append(uyari("Bilgi", "Grubun satışı yok", s["grup"], f"Fire {s['fire']:,.2f} ama grup dosyasında bu ürün grubu yok; oran hesaplanmadı."))
        if s["kayip_oran"] is not None and s["kayip_oran"] > esik:
            uy.append(uyari("Yüksek" if s["kayip_oran"] > esik * 2 else "Orta", "Kayıp oranı yüksek", s["grup"], f"Satış {s['satis']:,.2f}, kayıtlı fire {s['fire']:,.2f}, bilinmeyen kayıp {s['bilinmeyen'] or 0:,.2f}; "
                            f"kayıp oranı %{s['kayip_oran']} (eşik %{esik})."))
        if s["bilinmeyen_pay"] is not None and s["bilinmeyen_pay"] > 50:
            uy.append(uyari("Bilgi", "Kaybın çoğu kayıt dışı", s["grup"], f"Toplam kaybın %{s['bilinmeyen_pay']} kadarı sayım farkından ({s['bilinmeyen']:,.2f}). Hırsızlık, kayıtsız imha, teslim alma ve kasa hatalarını inceleyin."))
    grup_listesi = sorted(g.values(), key=lambda s: (s["kayip_oran"] is None, -(s["kayip_oran"] or 0), -s["fire"], katla(s["grup"])))
    ng = {}
    for f in fire:
        n = ng.setdefault(katla(f["neden"]), {"ad": f["neden"], "kayit": 0, "tutar": SIFIR})
        n["kayit"], n["tutar"] = n["kayit"] + 1, n["tutar"] + f["tutar"]
    nedenler, kum = sorted(ng.values(), key=lambda n: (-n["tutar"], katla(n["ad"]))), SIFIR
    for n in nedenler:
        n["pay"], n["sinif"] = _oran(n["tutar"], top_fire), "A" if top_fire and _oran(kum, top_fire) < 80 else "B"
        kum += n["tutar"]
        n["kum"] = _oran(kum, top_fire)
        if n["ad"] == YOK:
            uy.append(uyari("Orta", "Nedeni yazılmamış fire", YOK, f"{n['kayit']} kayıt, {n['tutar']:,.2f} (%{n['pay']})."))
    matris = [{"grup": s["grup"], "hucre": [sum((f["tutar"] for f in fire if katla(f["grup"]) == katla(s["grup"]) and katla(f["neden"]) == katla(n["ad"])), SIFIR) for n in nedenler], "toplam": s["fire"]} for s in grup_listesi if s["kayit"]]
    satis = sum((s["satis"] for s in g.values() if s["satis"] is not None), SIFIR)
    bilinmeyen = sum((s["bilinmeyen"] for s in g.values() if s["bilinmeyen"] is not None), SIFIR)
    return {"fire": fire, "gruplar": grup_listesi, "nedenler": nedenler, "matris": matris,
            "ozet": {"kayit": len(fire), "satis": satis if gruplar else None, "fire": top_fire, "bilinmeyen": bilinmeyen if gruplar else None, "toplam": top_fire + bilinmeyen, "fire_oran": _oran(top_fire, satis),
                     "kayip_oran": _oran(top_fire + bilinmeyen, satis)}, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, esik: Decimal) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ürün Grupları"
    _baslik(ws, ["Ürün Grubu", "Satış", "Kayıtlı Fire", "Fire %", "Sayım Farkı", "Bilinmeyen Kayıp", "Toplam Kayıp", "Kayıp %", "Kayıt Dışı Payı %", "Fire Kaydı", "Fire İçindeki Payı %"], (26, 16, 15, 9, 14, 15, 15, 9, 11, 9, 11), "B2")
    for s in r["gruplar"]:
        _satir(ws, [s["grup"], s["satis"], s["fire"], s["fire_oran"], s["fark"], s["bilinmeyen"], s["toplam"], s["kayip_oran"], s["bilinmeyen_pay"], s["kayit"] or None, s["fire_pay"]], para_s=range(2, 12),
               dolgu={8: None if s["kayip_oran"] is None else KIRMIZI if s["kayip_oran"] > esik * 2 else SARI if s["kayip_oran"] > esik else YESIL})
    o = r["ozet"]
    n = _satir(ws, ["TOPLAM", o["satis"], o["fire"], o["fire_oran"], -o["bilinmeyen"] if o["bilinmeyen"] is not None else None, o["bilinmeyen"], o["toplam"], o["kayip_oran"], None, o["kayit"], None], para_s=range(2, 9))
    for h in ws[n]:
        h.font = Font(bold=True)
    ws.append([])
    ws.append([f"Kayıp oranı eşiği %{esik}. Sayım farkı eksi ise eksik (kayıp), artı ise fazla."])
    ws = wb.create_sheet("Neden Pareto")
    _baslik(ws, ["Sıra", "Fire Nedeni", "Sınıf", "Kayıt", "Tutar", "Pay %", "Kümülatif %"], (6, 34, 7, 8, 16, 9, 12))
    for i, x in enumerate(r["nedenler"], 1):
        _satir(ws, [i, x["ad"], x["sinif"], x["kayit"], x["tutar"], x["pay"], x["kum"]], para_s=(5, 6, 7), dolgu={3: KIRMIZI if x["sinif"] == "A" else None})
    ws = wb.create_sheet("Grup × Neden")
    _baslik(ws, ["Ürün Grubu \\ Neden"] + [x["ad"] for x in r["nedenler"]] + ["Toplam"], (26,) + (16,) * (len(r["nedenler"]) + 1), "B2")
    for m in r["matris"]:
        _satir(ws, [m["grup"]] + [h or None for h in m["hucre"]] + [m["toplam"]], para_s=range(2, len(m["hucre"]) + 3))
    ws = wb.create_sheet("Fire Kayıtları")
    _baslik(ws, ["Tarih", "Ürün Grubu", "Neden", "Adet", "Tutar"], (12, 26, 34, 10, 16))
    for f in sorted(r["fire"], key=lambda f: (f["tarih"] or date.min, katla(f["grup"]))):
        _satir(ws, [f["tarih"], f["grup"], f["neden"], f["adet"], f["tutar"]], para_s=(5,), tarih_s=(1,))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(f_yol: Path, cikti: Path, g_yol: Path | None = None, esik: Decimal = Decimal(2)) -> dict:
    fire, uy1 = oku_fire(f_yol)
    gruplar, uy2 = oku_gruplar(g_yol) if g_yol else ([], [])
    r = analiz(fire, gruplar, esik)
    r["uyarilar"] = sirala(uy1 + uy2 + r["uyarilar"])
    rapor_yaz(cikti, r, esik)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Fire kayıtlarını ürün grubu ve neden bazında analiz eder, sayım farklarıyla birlikte kayıp oranını hesaplar.")
    p.add_argument("--fire", type=Path, help="Tarih, Ürün Grubu, Neden, Adet, Tutar")
    p.add_argument("--gruplar", type=Path, help="İsteğe bağlı: Ürün Grubu, Satış, Sayım Farkı; verilirse fire ve kayıp oranları hesaplanır")
    p.add_argument("--esik", default="2", help="Kayıp oranı bu yüzdeyi aşan ürün grubu uyarılır (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "fire_ve_kayip_analizi.xlsx")
    a = p.parse_args(argv)
    ornek = a.fire is None
    f_yol, g_yol = (ORNEK / "fire_kayitlari.csv", a.gruplar or ORNEK / "urun_gruplari.csv") if ornek else (a.fire, a.gruplar)
    esik = para(a.esik)
    if esik is None or esik <= 0:
        print("[X] --esik sıfırdan büyük bir yüzde olmalı")
        return 2
    for y in (f_yol, g_yol):
        if y is not None and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(f_yol, a.cikti, g_yol, esik)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {o['kayit']} fire kaydı · kayıtlı fire {o['fire']:,.2f}" + (f" (%{o['fire_oran']})" if o["fire_oran"] is not None else "") + (f" · bilinmeyen kayıp {o['bilinmeyen']:,.2f}" if o["bilinmeyen"] is not None else "")
          + (f" · toplam kayıp {o['toplam']:,.2f} (%{o['kayip_oran']})" if o["kayip_oran"] is not None else ""))
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
