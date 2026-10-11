import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
T = date(2026, 10, 5)


def ur(vardiya="1", hat="Hat 1", hedef=4000, uretilen=3900, fire=78, sure=480, t=T, urun="Kapak A"):
    return {"tarih": t, "vardiya": vardiya, "hat": hat, "urun": urun, "hedef": None if hedef is None else D(hedef), "uretilen": D(uretilen), "fire": D(fire), "sure": None if sure is None else D(sure)}


def du(vardiya="1", hat="Hat 1", neden="Arıza", sure=15, planli=False, t=T):
    return {"tarih": t, "vardiya": vardiya, "hat": hat, "neden": neden, "planli": planli, "sure": D(sure)}


class RaporTesti(unittest.TestCase):
    def test_hat_ve_vardiya(self):
        r = main.raporla([ur(), ur(hat="Hat 2", hedef=2400, uretilen=2450, fire=25, urun="Gövde B")], [du(neden="Kalıp değişimi", sure=30, planli=True), du(neden="Malzeme bekleme")])
        v = r["vardiyalar"][0]
        h1, h2 = v["hatlar"]
        self.assertEqual((h1["saglam"], h1["gerceklesme"], h1["fire_oran"], h1["planli"], h1["plansiz"], h1["kullanilabilirlik"]), (D(3822), D("95.55"), D("2.00"), D(30), D(15), D("96.67")))
        self.assertEqual((h2["gerceklesme"], h2["fire_oran"], h2["kullanilabilirlik"], h2["isaret"]), (D("101.04"), D("1.02"), D("100.00"), []))
        self.assertEqual((v["hedef"], v["uretilen"], v["saglam"], v["gerceklesme"], v["fire_oran"], v["kullanilabilirlik"]), (D(6400), D(6350), D(6247), D("97.61"), D("1.62"), D("98.39")))
        self.assertEqual(v["pareto"], [("Malzeme bekleme", False, D(15)), ("Kalıp değişimi", True, D(30))])

    def test_isaretler_ve_esikler(self):
        i = lambda **kw: main.raporla([ur(**kw)], [])["vardiyalar"][0]["isaret"]   # noqa: E731
        self.assertEqual(i(hedef=1000, uretilen=900, fire=0), [])                                  # tam %90 işaretlenmez
        self.assertEqual(i(hedef=1000, uretilen=899, fire=0), ["Hedefin altında"])
        self.assertEqual(i(hedef=1000, uretilen=1000, fire=30), [])                                # tam %3 işaretlenmez
        self.assertEqual(i(hedef=1000, uretilen=1000, fire=31), ["Fire yüksek"])
        self.assertEqual(i(hedef=1000, uretilen=800, fire=80), ["Hedefin altında", "Fire yüksek"])
        self.assertEqual(main.raporla([ur(hedef=1000, uretilen=950, fire=10)], [], hedef_esik=D(95), fire_esik=D(1))["vardiyalar"][0]["isaret"], ["Hedefin altında", "Fire yüksek"])
        v = main.raporla([ur(hedef=None, sure=None, uretilen=0, fire=0)], [du()])["vardiyalar"][0]
        self.assertEqual((v["hedef"], v["gerceklesme"], v["fire_oran"], v["kullanilabilirlik"], v["plansiz"]), (None, None, None, None, D(15)))

    def test_gun_toplami_ve_pareto(self):
        u = [ur("1"), ur("2", uretilen=3200, fire=160), ur("1", t=date(2026, 10, 6), uretilen=4000, fire=0)]
        d = [du("2", neden="Arıza", sure=75), du("2", neden="Ayar", sure=20), du("1", neden="arıza", sure=5), du("1", neden="Mola", sure=30, planli=True), du("3", neden="Yok", sure=9),
             du("1", hat="Hat 9", neden="Başka hat", sure=4)]
        r = main.raporla(u, d)
        self.assertEqual([(v["tarih"].day, v["vardiya"], v["plansiz"]) for v in r["vardiyalar"]], [(5, "1", D(9)), (5, "2", D(95)), (6, "1", D(0))])
        g = r["gunler"][0]
        self.assertEqual((g["hedef"], g["uretilen"], g["fire"], g["saglam"], g["gerceklesme"], g["fire_oran"], g["plansiz"], g["planli"]),
                         (D(8000), D(7100), D(238), D(6862), D("85.78"), D("3.35"), D(104), D(30)))
        self.assertEqual(g["kullanilabilirlik"], D("88.82"))                                        # (960 − 30 − 104) ÷ 930
        self.assertEqual(r["pareto"][0], ("Arıza", False, D(75)))
        self.assertEqual(sorted(u_["tur"] for u_ in r["uyarilar"]), ["Duruşun vardiyası üretim dosyasında yok", "Duruşun üretim satırı yok"])
        self.assertEqual(r["vardiyalar"][1]["hatlar"][0]["isaret"], ["Hedefin altında", "Fire yüksek"])

    def test_rapor_metni(self):
        m = main.rapor_metni(main.raporla([ur("2", uretilen=3200, fire=160)], [du("2", neden="Arıza - enjeksiyon", sure=75), du("2", neden="Ayar", sure=20)]))
        self.assertIn("=== 05.10.2026 · Vardiya 2 ===", m)
        self.assertIn("Üretim : 3.200 adet (hedef 4.000) · sağlam 3.040 · gerçekleşme %76", m)
        self.assertIn("Fire   : 160 adet (%5)", m)
        self.assertIn("Duruş  : plansız 95 dk, planlı 0 dk · kullanılabilirlik %80,21", m)
        self.assertIn("  - Hat 1 · Kapak A: sağlam 3.040 / hedef 4.000 (%76) · fire %5 · plansız duruş 95 dk [Hedefin altında, Fire yüksek]", m)
        self.assertIn("Plansız duruş nedenleri: Arıza - enjeksiyon 75 dk; Ayar 20 dk", m)
        self.assertEqual([main.sayi(x) for x in (None, D("1234.5"), D(1200), D("0.00"))], ["—", "1.234,50", "1.200", "0"])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "v.xlsx"
            r = main.calistir(ORNEK / "uretim.csv", ORNEK / "duruslar.csv", cikti)
            self.assertEqual([(v["vardiya"], v["saglam"], v["gerceklesme"], v["fire_oran"], v["plansiz"], v["kullanilabilirlik"], v["isaret"]) for v in r["vardiyalar"]],
                             [("1", D(6247), D("97.61"), D("1.62"), D(15), D("98.39"), []), ("2", D(5294), D("82.72"), D("3.75"), D(105), D("89.06"), ["Hedefin altında", "Fire yüksek"]),
                              ("3", D(5036), D("78.69"), D("3.15"), D(150), D("83.33"), ["Hedefin altında", "Fire yüksek"])])
            g = r["gunler"][0]
            self.assertEqual((g["hedef"], g["uretilen"], g["fire"], g["saglam"], g["gerceklesme"], g["fire_oran"]), (D(19200), D(17050), D(473), D(16577), D("86.34"), D("2.77")))
            h = r["vardiyalar"][2]["hatlar"][1]
            self.assertEqual((h["hat"], h["gerceklesme"], h["fire_oran"], h["kullanilabilirlik"]), ("Hat 2", D("58.75"), D("6.00"), D("66.67")))
            self.assertEqual([(n, s) for n, p, s in r["pareto"] if not p], [("Arıza - robot", D(150)), ("Arıza - enjeksiyon ünitesi", D(75)), ("Malzeme bekleme", D(25)), ("Ayar", D(20))])
            self.assertEqual(r["uyarilar"], [])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Vardiya Özeti", "Hat Bazında", "Duruş Nedenleri", "Vardiya Duruşları", "Uyarılar"])
            self.assertIn("Vardiya amiri notu:", (Path(tmp) / "vardiya_raporu.txt").read_text(encoding="utf-8"))

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            u, d = Path(tmp) / "u.csv", Path(tmp) / "d.csv"
            u.write_text("Tarih;Vardiya;Hat;Hedef Adet;Üretilen Adet;Fire Adet\n05.10.2026;A;H1;1.000;900;950\ndün;A;H1;1;1;0\n05.10.2026;B;H1;çok;yok;\n05.10.2026;B;H1;bin;500;\n06.10.2026;A;;;10;\n",
                         encoding="utf-8")
            d.write_text("Tarih;Vardiya;Hat;Duruş Nedeni;Tür;Süre (dk)\n05.10.2026;A;H1;Arıza;PLANSIZ;10\n05.10.2026;A;H1;Mola;Planlı duruş;20\n05.10.2026;A;H1;;;uzun\n", encoding="utf-8")
            ul, dl, uy = main.oku(u, d)
            self.assertEqual([(x["vardiya"], x["hat"], x["hedef"], x["uretilen"]) for x in ul], [("A", "H1", D(1000), D(900)), ("B", "H1", None, D(500)), ("A", "Belirtilmemiş", None, D(10))])
            self.assertEqual([(x["neden"], x["planli"], x["sure"]) for x in dl], [("Arıza", False, D(10)), ("Mola", True, D(20))])
            self.assertEqual(sorted(x["tur"] for x in uy), ["Değer okunamadı", "Duruş satırı okunamadı", "Fire üretilenden büyük", "Tarih veya vardiya okunamadı", "Üretilen adet okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--uretim", str(u), "--duruslar", str(d), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "05.10.2026", "--hedef-esik", "95", "--fire-esik", "2"]), 0)
            self.assertEqual(main.main(["--uretim", str(u), "--cikti", c]), 0)                       # duruş dosyası isteğe bağlı
            self.assertEqual(main.main(arg + ["--tarih", "dün"]), 2)
            self.assertEqual(main.main(arg + ["--tarih", "01.01.2020"]), 1)
            self.assertEqual(main.main(["--uretim", str(Path(tmp) / "yok.csv")]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
