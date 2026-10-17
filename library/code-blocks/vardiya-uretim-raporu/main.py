"""
Vardiya Üretim Raporu — Gerisi Bende kod bloğu
Üretim › Vardiya Amiri

Vardiya bazında üretim, hedef, duruş ve fire kayıtlarından vardiya raporu üretir.
  Sağlam adet         = üretilen adet − fire adet
  Hedef gerçekleşme % = sağlam adet ÷ hedef adet × 100
  Fire %              = fire adet ÷ üretilen adet × 100
  Kullanılabilirlik % = (planlı süre − planlı duruş − plansız duruş) ÷ (planlı süre − planlı duruş) × 100     (planlı süre girildiyse)
  İşaretler           = gerçekleşme < --hedef-esik (90) "Hedefin altında";  fire % > --fire-esik (3) "Fire yüksek"
Kod OEE hesaplamaz (performans için ideal çevrim süresi gerekir); yalnız girilen adet ve süreleri toplar. Eşikler işletmenize göre değişir.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


URETIM_SUTUN = {"tarih": ("tarih",), "vardiya": ("vardiya",), "hat": ("hat", "makine", "hat makine", "is merkezi"), "urun": ("urun", "urun adi", "stok kodu"),
                "hedef": ("hedef adet", "hedef", "planlanan adet"), "uretilen": ("uretilen adet", "uretilen", "uretim adedi", "toplam uretim"), "fire": ("fire adet", "fire", "hurda", "iskarta"),
                "sure": ("planli sure", "planli sure dk", "vardiya suresi", "sure dk")}
DURUS_SUTUN = {"tarih": ("tarih",), "vardiya": ("vardiya",), "hat": ("hat", "makine", "hat makine", "is merkezi"), "neden": ("durus nedeni", "neden", "durus"), "tur": ("tur", "durus turu"),
               "sure": ("sure dk", "sure", "dakika", "durus suresi")}
Z = Decimal(0)


def yuzde(a: Decimal, b: Decimal) -> Decimal | None:
    return yuvarla(a / b * 100) if b else None


def oku(uretim_yol: Path, durus_yol: Path | None) -> tuple[list[dict], list[dict], list[dict]]:
    uy, uretim, durus = [], [], []
    for r in kayitlar(uretim_yol, URETIM_SUTUN, ("tarih", "vardiya", "uretilen")):
        if not metin(r["vardiya"]) and not metin(r["tarih"]):
            continue
        kim = f"{uretim_yol.name} satır {r['_satir']}"
        t, ur, fire, hedef, sure = tarih(r["tarih"]), para(r["uretilen"]), para(r["fire"]), para(r["hedef"]), para(r["sure"])
        if not t or not metin(r["vardiya"]):
            uy.append(uyari("Yüksek", "Tarih veya vardiya okunamadı", kim, f"Tarih '{metin(r['tarih'])}', vardiya '{metin(r['vardiya'])}'. Satır rapora alınmadı."))
            continue
        if ur is None or ur < 0:
            uy.append(uyari("Yüksek", "Üretilen adet okunamadı", kim, f"'{metin(r['uretilen'])}'. Satır rapora alınmadı."))
            continue
        for v, ham, ad in ((fire, r["fire"], "Fire adet"), (hedef, r["hedef"], "Hedef adet"), (sure, r["sure"], "Planlı süre")):
            if v is None and metin(ham):
                uy.append(uyari("Orta", "Değer okunamadı", kim, f"{ad} '{metin(ham)}'. Boş sayıldı."))
        fire = fire or Z
        if fire > ur:
            uy.append(uyari("Yüksek", "Fire üretilenden büyük", kim, f"Fire {fire:f} > üretilen {ur:f}. Üretilen adet toplam (sağlam + fire) olmalı; satırı kontrol edin."))
        uretim.append({"tarih": t, "vardiya": metin(r["vardiya"]), "hat": metin(r["hat"]) or "Belirtilmemiş", "urun": metin(r["urun"]), "hedef": hedef, "uretilen": ur, "fire": fire, "sure": sure})
    if durus_yol:
        for r in kayitlar(durus_yol, DURUS_SUTUN, ("tarih", "vardiya", "sure")):
            if not metin(r["vardiya"]) and not metin(r["tarih"]):
                continue
            kim = f"{durus_yol.name} satır {r['_satir']}"
            t, sure = tarih(r["tarih"]), para(r["sure"])
            if not t or not metin(r["vardiya"]) or sure is None or sure < 0:
                uy.append(uyari("Yüksek", "Duruş satırı okunamadı", kim, f"Tarih '{metin(r['tarih'])}', vardiya '{metin(r['vardiya'])}', süre '{metin(r['sure'])}'. Satır alınmadı."))
                continue
            durus.append({"tarih": t, "vardiya": metin(r["vardiya"]), "hat": metin(r["hat"]) or "Belirtilmemiş", "neden": metin(r["neden"]) or "Belirtilmemiş",
                          "planli": katla(r["tur"]).startswith("planli"), "sure": sure})
    return uretim, durus, uy


def _ozet(satirlar: list[dict], duruslar: list[dict], hedef_esik: Decimal, fire_esik: Decimal) -> dict:
    hedefler = [s["hedef"] for s in satirlar if s["hedef"] is not None]
    sureler = [s["sure"] for s in satirlar if s["sure"] is not None]
    x = {"hedef": sum(hedefler, Z) if len(hedefler) == len(satirlar) and satirlar else None, "uretilen": sum((s["uretilen"] for s in satirlar), Z), "fire": sum((s["fire"] for s in satirlar), Z),
         "sure": sum(sureler, Z) if len(sureler) == len(satirlar) and satirlar else None, "planli": sum((d["sure"] for d in duruslar if d["planli"]), Z),
         "plansiz": sum((d["sure"] for d in duruslar if not d["planli"]), Z)}
    x["saglam"] = x["uretilen"] - x["fire"]
    x["gerceklesme"] = yuzde(x["saglam"], x["hedef"]) if x["hedef"] else None
    x["fire_oran"] = yuzde(x["fire"], x["uretilen"])
    net = x["sure"] - x["planli"] if x["sure"] is not None else None
    x["kullanilabilirlik"] = yuzde(max(net - x["plansiz"], Z), net) if net and net > 0 else None
    x["isaret"] = [m for m, k in (("Hedefin altında", x["gerceklesme"] is not None and x["gerceklesme"] < hedef_esik), ("Fire yüksek", x["fire_oran"] is not None and x["fire_oran"] > fire_esik)) if k]
    return x


def _pareto(duruslar: list[dict]) -> list[tuple]:
    g = defaultdict(lambda: Z)
    for d in duruslar:
        g[(d["neden"], d["planli"])] += d["sure"]
    return sorted(((n, p, s) for (n, p), s in g.items()), key=lambda x: (x[1], -x[2], x[0]))


def raporla(uretim: list[dict], durus: list[dict], hedef_esik: Decimal = Decimal(90), fire_esik: Decimal = Decimal(3)) -> dict:
    uy, vardiyalar = [], []
    anahtar = lambda x: (x["tarih"], katla(x["vardiya"]))   # noqa: E731
    for k in sorted({anahtar(u) for u in uretim}):
        us, ds = [u for u in uretim if anahtar(u) == k], [d for d in durus if anahtar(d) == k]
        v = {"tarih": k[0], "vardiya": us[0]["vardiya"], "hatlar": [], "pareto": _pareto(ds)} | _ozet(us, ds, hedef_esik, fire_esik)
        for hat in dict.fromkeys(katla(u["hat"]) for u in us):
            hs = [u for u in us if katla(u["hat"]) == hat]
            v["hatlar"].append({"hat": hs[0]["hat"], "urun": ", ".join(dict.fromkeys(u["urun"] for u in hs if u["urun"]))} | _ozet(hs, [d for d in ds if katla(d["hat"]) == hat], hedef_esik, fire_esik))
        hatlar = {katla(u["hat"]) for u in us}
        for d in ds:
            if katla(d["hat"]) not in hatlar:
                uy.append(uyari("Orta", "Duruşun üretim satırı yok", f"{k[0]:%d.%m.%Y} · {d['vardiya']} · {d['hat']}", f"'{d['neden']}' duruşu ({d['sure']:f} dk) vardiya toplamına girdi ama hat tablosunda görünmez."))
        vardiyalar.append(v)
    vk = {anahtar(u) for u in uretim}
    for d in durus:
        if anahtar(d) not in vk:
            uy.append(uyari("Yüksek", "Duruşun vardiyası üretim dosyasında yok", f"{d['tarih']:%d.%m.%Y} · {d['vardiya']}", f"'{d['neden']}' duruşu ({d['sure']:f} dk) hiçbir vardiyaya eklenemedi."))
    gunler = []
    for t in sorted({u["tarih"] for u in uretim}):
        gunler.append({"tarih": t} | _ozet([u for u in uretim if u["tarih"] == t], [d for d in durus if d["tarih"] == t and anahtar(d) in vk], hedef_esik, fire_esik))
    return {"vardiyalar": vardiyalar, "gunler": gunler, "pareto": _pareto([d for d in durus if anahtar(d) in vk]), "uyarilar": uy, "esik": (hedef_esik, fire_esik)}


def sayi(x) -> str:
    if x is None:
        return "—"
    x = Decimal(x)
    s = f"{x:,.2f}" if x != x.to_integral_value() else f"{int(x):,}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def rapor_metni(r: dict) -> str:
    b = []
    for v in r["vardiyalar"]:
        b.append(f"=== {v['tarih']:%d.%m.%Y} · Vardiya {v['vardiya']} ===")
        b.append(f"Üretim : {sayi(v['uretilen'])} adet (hedef {sayi(v['hedef'])}) · sağlam {sayi(v['saglam'])} · gerçekleşme %{sayi(v['gerceklesme'])}")
        b.append(f"Fire   : {sayi(v['fire'])} adet (%{sayi(v['fire_oran'])})")
        b.append(f"Duruş  : plansız {sayi(v['plansiz'])} dk, planlı {sayi(v['planli'])} dk" + (f" · kullanılabilirlik %{sayi(v['kullanilabilirlik'])}" if v["kullanilabilirlik"] is not None else ""))
        b.append("Hat bazında:")
        for h in v["hatlar"]:
            b.append(f"  - {h['hat']}" + (f" · {h['urun']}" if h["urun"] else "") + f": sağlam {sayi(h['saglam'])} / hedef {sayi(h['hedef'])} (%{sayi(h['gerceklesme'])}) · fire %{sayi(h['fire_oran'])}"
                     f" · plansız duruş {sayi(h['plansiz'])} dk" + (f" [{', '.join(h['isaret'])}]" if h["isaret"] else ""))
        plansiz = [f"{n} {sayi(s)} dk" for n, p, s in v["pareto"] if not p]
        if plansiz:
            b.append("Plansız duruş nedenleri: " + "; ".join(plansiz))
        b += ["Vardiya amiri notu: ", ""]
    return "\n".join(b)


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Vardiya Özeti"
    bas = ["Hedef Adet", "Üretilen Adet", "Fire Adet", "Sağlam Adet", "Gerçekleşme %", "Fire %", "Planlı Duruş (dk)", "Plansız Duruş (dk)", "Kullanılabilirlik %", "İşaret"]
    _baslik(ws, ["Tarih", "Vardiya", *bas], (12, 14, 12, 12, 10, 12, 12, 9, 12, 12, 13, 30))

    def yaz(ws, on, x, kalin=False):
        n = _satir(ws, [*on, x["hedef"], x["uretilen"], x["fire"], x["saglam"], x["gerceklesme"], x["fire_oran"], x["planli"] or None, x["plansiz"] or None, x["kullanilabilirlik"], ", ".join(x["isaret"])],
                   para_s=(len(on) + 5, len(on) + 6, len(on) + 9), tarih_s=(1,),
                   dolgu={len(on) + 5: KIRMIZI if "Hedefin altında" in x["isaret"] else None, len(on) + 6: KIRMIZI if "Fire yüksek" in x["isaret"] else None})
        if kalin:
            for c in ws[n]:
                c.font = Font(bold=True)

    for g in r["gunler"]:
        for v in (v for v in r["vardiyalar"] if v["tarih"] == g["tarih"]):
            yaz(ws, [v["tarih"], v["vardiya"]], v)
        yaz(ws, [g["tarih"], "Gün toplamı"], g, True)
    ws.append([])
    ws.append([f"Gerçekleşme = sağlam adet ÷ hedef. Eşikler: gerçekleşme < %{r['esik'][0]} hedefin altında, fire > %{r['esik'][1]} yüksek. OEE hesaplanmaz."])
    ws = wb.create_sheet("Hat Bazında")
    _baslik(ws, ["Tarih", "Vardiya", "Hat", "Ürün", *bas], (12, 14, 16, 22, 12, 12, 10, 12, 12, 9, 12, 12, 13, 30))
    for v in r["vardiyalar"]:
        for h in v["hatlar"]:
            yaz(ws, [v["tarih"], v["vardiya"], h["hat"], h["urun"]], h)
    ws = wb.create_sheet("Duruş Nedenleri")
    _baslik(ws, ["Duruş Nedeni", "Tür", "Toplam Süre (dk)", "Pay % (türü içinde)", "Birikimli %"], (36, 12, 14, 14, 12))
    for planli in (False, True):
        l = [x for x in r["pareto"] if x[1] == planli]
        toplam, bir = sum((x[2] for x in l), Z), Z
        for n, p, s in l:
            bir += s
            _satir(ws, [n, "Planlı" if p else "Plansız", s, yuzde(s, toplam), yuzde(bir, toplam)], para_s=(4, 5), dolgu={2: None if p else SARI})
    ws = wb.create_sheet("Vardiya Duruşları")
    _baslik(ws, ["Tarih", "Vardiya", "Duruş Nedeni", "Tür", "Süre (dk)"], (12, 14, 36, 12, 12))
    for v in r["vardiyalar"]:
        for n, p, s in v["pareto"]:
            _satir(ws, [v["tarih"], v["vardiya"], n, "Planlı" if p else "Plansız", s], tarih_s=(1,))
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)
    (cikti.parent / "vardiya_raporu.txt").write_text(rapor_metni(r), encoding="utf-8")


def calistir(uretim_yol: Path, durus_yol: Path | None, cikti: Path, gun: date | None = None, hedef_esik: Decimal = Decimal(90), fire_esik: Decimal = Decimal(3)) -> dict:
    uretim, durus, uy = oku(uretim_yol, durus_yol)
    if gun:
        uretim, durus = [u for u in uretim if u["tarih"] == gun], [d for d in durus if d["tarih"] == gun]
    if not uretim:
        raise ValueError(f"{uretim_yol.name}: raporlanacak üretim satırı yok" + (f" ({gun:%d.%m.%Y})." if gun else "."))
    r = raporla(uretim, durus, hedef_esik, fire_esik)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, r["uyarilar"])
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Vardiya bazında üretim, hedef, duruş ve fire kayıtlarından vardiya raporu üretir.")
    p.add_argument("--uretim", type=Path, default=ORNEK / "uretim.csv", help="Tarih, Vardiya, Hat, Ürün, Hedef Adet, Üretilen Adet, Fire Adet, Planlı Süre (dk)")
    p.add_argument("--duruslar", type=Path, help="Tarih, Vardiya, Hat, Duruş Nedeni, Tür (Planlı / Plansız), Süre (dk)")
    p.add_argument("--tarih", help="Yalnız bu günü raporla (GG.AA.YYYY); verilmezse dosyadaki tüm günler")
    p.add_argument("--hedef-esik", default="90", help="Bu gerçekleşme yüzdesinin altı işaretlenir (varsayılan 90)")
    p.add_argument("--fire-esik", default="3", help="Bu fire yüzdesinin üstü işaretlenir (varsayılan 3)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "vardiya_uretim_raporu.xlsx")
    a = p.parse_args(argv)
    durus = a.duruslar or (ORNEK / "duruslar.csv" if a.uretim == ORNEK / "uretim.csv" else None)
    for y in (a.uretim, durus):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    gun, he, fe = tarih(a.tarih) if a.tarih else None, para(a.hedef_esik), para(a.fire_esik)
    if (a.tarih and not gun) or he is None or fe is None or he < 0 or fe < 0:
        print("[X] --tarih GG.AA.YYYY biçiminde; --hedef-esik ve --fire-esik sıfır veya pozitif sayı olmalı")
        return 2
    try:
        r = calistir(a.uretim, durus, a.cikti, gun, he, fe)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    for v in r["vardiyalar"]:
        print(f"[OK] {v['tarih']:%d.%m.%Y} · Vardiya {v['vardiya']}: sağlam {sayi(v['saglam'])} / hedef {sayi(v['hedef'])} (%{sayi(v['gerceklesme'])}) · fire %{sayi(v['fire_oran'])} · "
              f"plansız duruş {sayi(v['plansiz'])} dk" + (f" [{', '.join(v['isaret'])}]" if v["isaret"] else ""))
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    print(f"[OK] Vardiya raporu metni: {(a.cikti.parent / 'vardiya_raporu.txt').resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
