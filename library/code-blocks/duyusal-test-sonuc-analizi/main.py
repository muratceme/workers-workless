"""
Duyusal Test Sonuç Analizi — Gerisi Bende kod bloğu
Ar-Ge ve Ürün Geliştirme › Ar-Ge Uzmanı
Panel değerlendirme puanlarını numune ve özellik bazında istatistiksel olarak karşılaştırır.
  Veri              = her satır bir panelistin bir numuneye verdiği puanlar; Panelist ve Numune dışındaki sayısal sütunlar özelliktir
  Varyans analizi   = her panelist her numuneyi bir kez puanladıysa iki yönlü ANOVA (numune + panelist, etkileşimsiz);
                      değilse tek yönlü ANOVA (numune)
  Anlamlılık        = numune etkisinin p değeri < --alfa (0,05)
  Çoklu karşılaştırma = Fisher LSD: |ortalama farkı| > t(1 − α/2; hata sd) × √(KO_hata × (1/nᵢ + 1/nⱼ))
  Harf grupları     = dengeli veride aynı harfi paylaşan numuneler arasında anlamlı fark yoktur
  Panelist notu     = panelistin özellik ortalaması panel ortalamasından --panelist (1,5) puandan fazla sapıyorsa
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SIFIR = Decimal(0)
PANELIST = ("panelist", "panelist no", "degerlendirici", "katilimci", "kisi")
NUMUNE = ("numune", "numune kodu", "ornek", "urun", "kod")
ATLA = ("oturum", "tarih", "tekrar", "sira", "aciklama", "not")


# ---- dağılım fonksiyonları (dış kütüphane gerektirmez) ----
def _betacf(a: float, b: float, x: float) -> float:
    qab, qap, qam = a + b, a + 1, a - 1
    c, d = 1.0, 1 - qab * x / qap
    d = 1 / (d if abs(d) > 1e-300 else 1e-300)
    h = d
    for m in range(1, 300):
        m2 = 2 * m
        aa = m * (b - m) * x / ((qam + m2) * (a + m2))
        d = 1 + aa * d
        d = 1 / (d if abs(d) > 1e-300 else 1e-300)
        c = 1 + aa / c if abs(1 + aa / c) > 1e-300 else 1e-300
        h *= d * c
        aa = -(a + m) * (qab + m) * x / ((a + m2) * (qap + m2))
        d = 1 + aa * d
        d = 1 / (d if abs(d) > 1e-300 else 1e-300)
        c = 1 + aa / c if abs(1 + aa / c) > 1e-300 else 1e-300
        dl = d * c
        h *= dl
        if abs(dl - 1) < 1e-12:
            break
    return h


def beta_ink(a: float, b: float, x: float) -> float:
    """Düzenlenmiş eksik beta fonksiyonu I_x(a, b)."""
    if x <= 0:
        return 0.0
    if x >= 1:
        return 1.0
    ln = math.lgamma(a + b) - math.lgamma(a) - math.lgamma(b) + a * math.log(x) + b * math.log(1 - x)
    if x < (a + 1) / (a + b + 2):
        return math.exp(ln) * _betacf(a, b, x) / a
    return 1 - math.exp(ln) * _betacf(b, a, 1 - x) / b


def f_p(f: float, d1: int, d2: int) -> float:
    return 1.0 if f <= 0 else beta_ink(d2 / 2, d1 / 2, d2 / (d2 + d1 * f))


def t_p(t: float, df: int) -> float:
    """İki yönlü p."""
    return beta_ink(df / 2, 0.5, df / (df + t * t))


def t_kritik(alfa: float, df: int) -> float:
    alt, ust = 0.0, 1000.0
    for _ in range(200):
        orta = (alt + ust) / 2
        if t_p(orta, df) > alfa:
            alt = orta
        else:
            ust = orta
    return (alt + ust) / 2


# ---- okuma ----
def oku(yol: Path) -> tuple[list[dict], list[str], list[dict]]:
    satirlar = tablo_oku(yol)
    for bi, r in enumerate(satirlar[:15]):
        b = [katla(c) for c in r]
        pi = next((b.index(a) for a in PANELIST if a in b), None)
        ni = next((b.index(a) for a in NUMUNE if a in b and b.index(a) != pi), None)
        if pi is not None and ni is not None:
            break
    else:
        raise ValueError(f"{yol.name}: başlık satırı bulunamadı; Panelist ve Numune sütunları gerekli")
    ozellik = [(j, metin(satirlar[bi][j])) for j in range(len(satirlar[bi])) if j not in (pi, ni) and metin(satirlar[bi][j]) and katla(satirlar[bi][j]) not in ATLA]
    uy, l = [], []
    for n, r in enumerate(satirlar[bi + 1:], bi + 2):
        p, nm = metin(r[pi] if pi < len(r) else ""), metin(r[ni] if ni < len(r) else "")
        if not p and not nm:
            continue
        puan = {}
        for j, ad in ozellik:
            v = r[j] if j < len(r) else None
            if metin(v):
                x = para(v)
                if x is None:
                    uy.append(uyari("Orta", "Puan okunamadı", f"{p} · {nm}", f"Satır {n}: {ad} '{metin(v)}' sayı değil; bu puan alınmadı."))
                    continue
                puan[ad] = float(x)
        if not p or not nm:
            uy.append(uyari("Yüksek", "Satır okunamadı", f"{p} · {nm}", f"Satır {n}: panelist ya da numune boş."))
            continue
        l.append({"panelist": p, "numune": nm, "puan": puan})
    ozellik_ad = [ad for _, ad in ozellik if any(ad in x["puan"] for x in l)]
    return l, ozellik_ad, uy


# ---- analiz ----
def _ort(l):
    return sum(l) / len(l)


def anova(veri: list[tuple[str, str, float]], alfa: float = 0.05) -> dict:
    """veri: (panelist, numune, puan)."""
    numuneler = sorted({n for _, n, _ in veri}, key=katla)
    panelistler = sorted({p for p, _, _ in veri}, key=katla)
    hucre = {}
    for p, n, x in veri:
        hucre.setdefault((p, n), []).append(x)
    dengeli = all(len(hucre.get((p, n), [])) == 1 for p in panelistler for n in numuneler) and len(panelistler) > 1
    tum = [x for _, _, x in veri]
    g = _ort(tum)
    grup = {n: [x for _, m, x in veri if m == n] for n in numuneler}
    kt_top = sum((x - g) ** 2 for x in tum)
    kt_num = sum(len(v) * (_ort(v) - g) ** 2 for v in grup.values())
    df_num = len(numuneler) - 1
    if dengeli:
        kt_pan = sum(len(numuneler) * (_ort([x for p, _, x in veri if p == q]) - g) ** 2 for q in panelistler)
        df_pan = len(panelistler) - 1
        kt_hata, df_hata = kt_top - kt_num - kt_pan, df_num * df_pan
    else:
        kt_pan, df_pan = None, None
        kt_hata, df_hata = kt_top - kt_num, len(tum) - len(numuneler)
    sonuc = {"numuneler": numuneler, "dengeli": dengeli, "ortalama": {n: _ort(v) for n, v in grup.items()}, "n": {n: len(v) for n, v in grup.items()},
             "sd": {n: (sum((x - _ort(v)) ** 2 for x in v) / (len(v) - 1)) ** 0.5 if len(v) > 1 else None for n, v in grup.items()},
             "df_num": df_num, "df_hata": df_hata, "f": None, "p": None, "anlamli": None, "lsd": None, "ciftler": [], "harf": {}}
    if df_num < 1 or df_hata < 1:
        return sonuc
    ko_hata = kt_hata / df_hata
    if ko_hata <= 0:
        sonuc |= {"f": None, "p": 0.0 if kt_num > 0 else 1.0, "anlamli": kt_num > 0}
        return sonuc
    f = (kt_num / df_num) / ko_hata
    p = f_p(f, df_num, df_hata)
    tk = t_kritik(alfa, df_hata)
    sonuc |= {"f": f, "p": p, "anlamli": p < alfa, "ko_hata": ko_hata}
    o = sonuc["ortalama"]
    for i, a in enumerate(numuneler):
        for b in numuneler[i + 1:]:
            lsd = tk * (ko_hata * (1 / len(grup[a]) + 1 / len(grup[b]))) ** 0.5
            fark = o[a] - o[b]
            sonuc["ciftler"].append({"a": a, "b": b, "fark": fark, "lsd": lsd, "anlamli": abs(fark) > lsd})
    if dengeli:
        lsd = tk * (2 * ko_hata / len(panelistler)) ** 0.5
        sonuc["lsd"] = lsd
        sirali = sorted(numuneler, key=lambda n: -o[n])
        gruplar = []
        for i in range(len(sirali)):
            j = i
            while j + 1 < len(sirali) and o[sirali[i]] - o[sirali[j + 1]] <= lsd:
                j += 1
            aralik = (i, j)
            if not any(g0 <= i and j <= g1 for g0, g1 in gruplar):
                gruplar.append(aralik)
        harf = {n: "" for n in numuneler}
        for k, (i, j) in enumerate(gruplar):
            for n in sirali[i:j + 1]:
                harf[n] += "abcdefghijklmnopqrstuvwxyz"[k]
        sonuc["harf"] = harf
    return sonuc


def analiz(kayit: list[dict], ozellikler: list[str], alfa: float = 0.05, panelist_esik: float = 1.5) -> dict:
    if not kayit or not ozellikler:
        raise ValueError("Değerlendirme puanı yok.")
    uy, sonuc = [], []
    for oz in ozellikler:
        veri = [(k["panelist"], k["numune"], k["puan"][oz]) for k in kayit if oz in k["puan"]]
        a = anova(veri, alfa)
        a["ozellik"] = oz
        if not a["dengeli"]:
            uy.append(uyari("Bilgi", "Dengesiz veri", oz, "Her panelist her numuneyi tam bir kez puanlamamış; panelist etkisi ayrılmadan tek yönlü ANOVA yapıldı."))
        if a["f"] is None and a["p"] is None:
            uy.append(uyari("Orta", "Analiz yapılamadı", oz, "En az iki numune ve yeterli tekrar gerekli."))
        panel_ort = _ort([x for _, _, x in veri])
        for p in sorted({p for p, _, _ in veri}, key=katla):
            po = _ort([x for q, _, x in veri if q == p])
            if abs(po - panel_ort) > panelist_esik:
                uy.append(uyari("Bilgi", "Panelist ortalamadan sapıyor", f"{oz} · {p}", f"Panelist ortalaması {po:.2f}, panel {panel_ort:.2f}. Ölçek kullanımı ya da eğitim ihtiyacı açısından inceleyin."))
        sonuc.append(a)
    numuneler = sorted({k["numune"] for k in kayit}, key=katla)
    return {"ozellikler": sonuc, "numuneler": numuneler, "panelist": len({katla(k["panelist"]) for k in kayit}), "uyarilar": uy}


def _y(x, hane=2):
    return None if x is None else round(x, hane)


def rapor_yaz(cikti: Path, r: dict, alfa: float) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    nm = r["numuneler"]
    _baslik(ws, ["Özellik"] + [f"{n} ort." for n in nm] + ["F", "p", "Numuneler arasında fark", "LSD", "Analiz"], (20,) + (12,) * len(nm) + (8, 8, 14, 8, 14), "B2")
    for a in r["ozellikler"]:
        hucre = [(f"{a['ortalama'][n]:.2f} {a['harf'].get(n, '')}".strip() if n in a["ortalama"] else None) for n in nm]
        _satir(ws, [a["ozellik"]] + hucre + [_y(a["f"]), _y(a["p"], 4), None if a["anlamli"] is None else "Anlamlı" if a["anlamli"] else "Anlamlı değil", _y(a["lsd"]),
                    "İki yönlü ANOVA" if a["dengeli"] else "Tek yönlü ANOVA"], dolgu={len(nm) + 4: YESIL if a["anlamli"] else None})
    ws.append([])
    ws.append([f"α = {alfa}. Aynı harfi paylaşan numuneler arasında anlamlı fark yoktur (Fisher LSD). Panelist sayısı {r['panelist']}."])
    ws = wb.create_sheet("Çift Karşılaştırma")
    _baslik(ws, ["Özellik", "Numune A", "Numune B", "Fark (A − B)", "LSD", "Sonuç"], (20, 12, 12, 11, 9, 14))
    for a in r["ozellikler"]:
        for c in a["ciftler"]:
            _satir(ws, [a["ozellik"], c["a"], c["b"], _y(c["fark"]), _y(c["lsd"]), "Anlamlı" if c["anlamli"] else "Anlamlı değil"], dolgu={6: YESIL if c["anlamli"] else None})
    ws = wb.create_sheet("Tanımlayıcı")
    _baslik(ws, ["Özellik", "Numune", "Puan Sayısı", "Ortalama", "Standart Sapma"], (20, 12, 9, 10, 12))
    for a in r["ozellikler"]:
        for n in a["numuneler"]:
            _satir(ws, [a["ozellik"], n, a["n"][n], _y(a["ortalama"][n]), _y(a["sd"][n])])
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, alfa: float = 0.05, panelist_esik: float = 1.5) -> dict:
    kayit, ozellikler, uy = oku(yol)
    r = analiz(kayit, ozellikler, alfa, panelist_esik)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, alfa)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Panel değerlendirme puanlarını numune ve özellik bazında istatistiksel olarak karşılaştırır.")
    p.add_argument("--puanlar", type=Path, help="Panelist, Numune ve her özellik için bir sütun (ör. Renk, Koku, Tat, Doku, Genel Beğeni)")
    p.add_argument("--alfa", default="0,05", help="Anlamlılık düzeyi (varsayılan 0,05)")
    p.add_argument("--panelist", default="1,5", help="Panelist ortalaması panelden bu kadar puan saparsa bilgi verilir (varsayılan 1,5)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "duyusal_test_analizi.xlsx")
    a = p.parse_args(argv)
    yol = a.puanlar or ORNEK / "panel_puanlari.csv"
    alfa, pe = para(a.alfa), para(a.panelist)
    if alfa is None or not 0 < alfa < 1 or pe is None or pe <= 0:
        print("[X] --alfa 0 ile 1 arası, --panelist sıfırdan büyük olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, float(alfa), float(pe))
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(r['numuneler'])} numune · {r['panelist']} panelist · {len(r['ozellikler'])} özellik")
    for o in r["ozellikler"]:
        sira = sorted(o["ortalama"], key=lambda n: -o["ortalama"][n])
        print(f"     {o['ozellik']}: " + ", ".join(f"{n} {o['ortalama'][n]:.2f}{o['harf'].get(n, '')}" for n in sira)
              + (f" · p = {o['p']:.4f} ({'anlamlı' if o['anlamli'] else 'anlamlı değil'})" if o["p"] is not None else ""))
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
