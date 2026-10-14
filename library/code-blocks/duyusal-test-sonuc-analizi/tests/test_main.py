import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
VERI = [("P1", "A", 7), ("P1", "B", 5), ("P1", "C", 3), ("P2", "A", 8), ("P2", "B", 6), ("P2", "C", 3), ("P3", "A", 6), ("P3", "B", 5), ("P3", "C", 4)]


class DagilimTesti(unittest.TestCase):
    def test_kritik_degerler(self):
        self.assertAlmostEqual(main.f_p(4.103, 2, 10), 0.05, places=3)
        self.assertAlmostEqual(main.f_p(18.0, 2, 4), 0.01, places=3)
        self.assertAlmostEqual(main.t_kritik(0.05, 10), 2.228, places=3)
        self.assertAlmostEqual(main.t_kritik(0.05, 4), 2.776, places=3)
        self.assertAlmostEqual(main.t_p(1.96, 10000), 0.05, places=3)
        self.assertEqual((main.f_p(0, 2, 4), main.beta_ink(2, 3, 0), main.beta_ink(2, 3, 1)), (1.0, 0.0, 1.0))
        self.assertAlmostEqual(main.beta_ink(1, 1, 0.3), 0.3, places=10)


class AnovaTesti(unittest.TestCase):
    def test_iki_yonlu(self):
        a = main.anova(VERI)
        self.assertTrue(a["dengeli"])
        self.assertAlmostEqual(a["f"], 16.5455, places=3)
        self.assertTrue(0.01 < a["p"] < 0.02 and a["anlamli"])
        self.assertAlmostEqual(a["lsd"], 1.7721, places=3)
        self.assertEqual((a["harf"], [(c["a"], c["b"], c["anlamli"]) for c in a["ciftler"]]), ({"A": "a", "B": "a", "C": "b"}, [("A", "B", False), ("A", "C", True), ("B", "C", True)]))
        self.assertEqual((a["df_num"], a["df_hata"]), (2, 4))

    def test_tek_yonlu_ve_analiz(self):
        a = main.anova(VERI[:-1])
        self.assertEqual((a["dengeli"], a["df_hata"], a["harf"], a["lsd"]), (False, 5, {}, None))
        self.assertTrue(a["anlamli"])
        k = [{"panelist": p, "numune": n, "puan": {"Tat": float(x), "Renk": 5.0}} for p, n, x in VERI] + [{"panelist": "P4", "numune": "A", "puan": {"Tat": 1.0}}]
        r = main.analiz(k, ["Tat", "Renk"])
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Dengesiz veri", "Tat"), ("Panelist ortalamadan sapıyor", "Tat · P4")])
        renk = next(o for o in r["ozellikler"] if o["ozellik"] == "Renk")
        self.assertEqual((renk["f"], renk["p"], renk["anlamli"]), (None, 1.0, False))
        self.assertEqual((r["numuneler"], r["panelist"]), (["A", "B", "C"], 4))
        with self.assertRaises(ValueError):
            main.analiz([], ["Tat"])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "panel_puanlari.csv", cikti)
            self.assertEqual([(o["ozellik"], "".join(o["harf"][n] + "/" for n in r["numuneler"]), o["anlamli"], round(o["p"], 4)) for o in r["ozellikler"]],
                             [("Görünüş", "a/a/a/", False, 0.9027), ("Koku", "a/ab/b/", True, 0.0498), ("Tat", "a/a/a/", False, 0.1438), ("Doku", "a/a/b/", True, 0.0027),
                              ("Genel Beğeni", "a/a/a/", False, 0.4222)])
            self.assertEqual([(u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Panelist ortalamadan sapıyor", "Doku · P12"), ("Panelist ortalamadan sapıyor", "Genel Beğeni · P12"), ("Panelist ortalamadan sapıyor", "Tat · P12")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Çift Karşılaştırma", "Tanımlayıcı", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "p.csv"
            y.write_text("Oturum;Değerlendirici;Örnek;Tat;Koku\n1;P1;A;7;5\n1;P1;B;x;6\n1;;C;3;3\n1;P2;A;8;\n", encoding="utf-8")
            l, oz, uy = main.oku(y)
            self.assertEqual(([(x["panelist"], x["numune"], x["puan"]) for x in l], oz, [u["tur"] for u in uy]),
                             ([("P1", "A", {"Tat": 7.0, "Koku": 5.0}), ("P1", "B", {"Koku": 6.0}), ("P2", "A", {"Tat": 8.0})], ["Tat", "Koku"], ["Puan okunamadı", "Satır okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--puanlar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--alfa", "0,01", "--panelist", "2"]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--alfa", "1"], ["--panelist", "0"])], [2] * 2)
            self.assertEqual(main.main(["--puanlar", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Ad;Puan\nA;5\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
