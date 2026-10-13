"""
Geri Çağırma (İzlenebilirlik) Simülasyonu — Workers / Workless kod bloğu
Kalite Güvence ve Gıda Güvenliği › Kalite Güvence Uzmanı

Bir hammadde lotundan (veya üretim partisinden) yola çıkıp etkilenen üretim partilerini, sevk edilen müşterileri ve miktarları çıkarır.
  İleri izleme (--lot)    = lotu kullanan üretim partileri; o parti başka partide yarı mamul olarak kullanıldıysa o partiler de (zincir boyunca)
  Geri izleme (--parti)   = partinin kullandığı hammadde lotları (yarı mamul üzerinden gelenler dahil) + partiyi yarı mamul olarak kullanan partiler
  Sevk edilen             = etkilenen partinin sevkiyat kayıtları toplamı
  İç kullanım             = partinin başka üretim partilerinde yarı mamul olarak kullanılan miktarı
  Depoda olması gereken   = üretilen − sevk edilen − iç kullanım                                    (negatifse kayıtlar tutarsız)
  Müşteri bazında         = müşteriye giden toplam miktar, parti ve sevkiyat sayısı, ilk ve son sevk tarihi
Yarı mamul bağlantısı, kullanım kaydındaki "Hammadde Lot" değerinin bir üretim partisi numarasına eşit olmasıyla kurulur.
Kod kayıtları birleştirir; depodaki fiili miktarı sayım doğrular. Gerçek geri çağırma kararı ve bildirimi yetkili kişilerce, yürürlükteki mevzuata göre yapılır.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


URETIM = {"parti": ("uretim partisi", "parti no", "parti", "uretim parti no", "lot no"), "urun": ("urun", "urun adi", "mamul"), "tarih": ("uretim tarihi", "tarih"),
          "miktar": ("uretilen miktar", "miktar", "uretim miktari"), "birim": ("birim",)}
KULLANIM = {"parti": URETIM["parti"], "hammadde": ("hammadde", "bilesen", "malzeme"), "lot": ("hammadde lot", "hammadde lot no", "lot", "girdi lot", "hammadde lotu"), "miktar": ("kullanilan miktar", "miktar")}
SEVK = {"parti": URETIM["parti"], "musteri": ("musteri", "alici", "cari"), "tarih": ("sevk tarihi", "tarih"), "miktar": ("miktar", "sevk miktari"), "irsaliye": ("irsaliye no", "irsaliye", "belge no")}
Z = Decimal(0)


def olcu(x) -> Decimal | None:
    """Miktar: 12,5 = 12.5; 1.250,5 = bin iki yüz elli buçuk. Tek başına nokta ondalık ayırıcıdır."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    else:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def oku(uretim_yol: Path, kullanim_yol: Path, sevk_yol: Path | None = None) -> tuple[dict, list[dict], list[dict], list[dict]]:
    uy, uretim, kullanim, sevk = [], {}, [], []
    for r in kayitlar(uretim_yol, URETIM, ("parti", "miktar")):
        p, m = metin(r["parti"]), olcu(r["miktar"])
        if not p:
            continue
        if m is None or m <= 0:
            uy.append(uyari("Yüksek", "Üretim kaydı okunamadı", f"{p} (satır {r['_satir']})", f"Üretilen miktar '{metin(r['miktar'])}'. Satır alınmadı."))
            continue
        if katla(p) in uretim:
            uy.append(uyari("Orta", "Mükerrer üretim partisi", p, f"Parti üretim dosyasında birden fazla satırda (satır {r['_satir']}). Miktarlar toplandı."))
            uretim[katla(p)]["miktar"] += m
            continue
        uretim[katla(p)] = {"parti": p, "urun": metin(r["urun"]), "tarih": tarih(r["tarih"]), "miktar": m, "birim": metin(r["birim"])}
    for r in kayitlar(kullanim_yol, KULLANIM, ("parti", "lot")):
        p, lot = metin(r["parti"]), metin(r["lot"])
        if not p and not lot:
            continue
        m = olcu(r["miktar"])
        if not p or not lot or (m is None and metin(r["miktar"])) or (m is not None and m < 0):
            uy.append(uyari("Yüksek", "Kullanım kaydı okunamadı", f"{p or '?'} · {lot or '?'} (satır {r['_satir']})", f"Parti, lot veya miktar ('{metin(r['miktar'])}') eksik ya da hatalı. Satır alınmadı; izleme zinciri eksik kalabilir."))
            continue
        kullanim.append({"parti": p, "parti_k": katla(p), "hammadde": metin(r["hammadde"]), "lot": lot, "lot_k": katla(lot), "miktar": m})
    if sevk_yol:
        for r in kayitlar(sevk_yol, SEVK, ("parti", "musteri", "miktar")):
            p, mu = metin(r["parti"]), metin(r["musteri"])
            if not p and not mu:
                continue
            m = olcu(r["miktar"])
            if not p or not mu or m is None or m <= 0:
                uy.append(uyari("Yüksek", "Sevkiyat kaydı okunamadı", f"{p or '?'} · {mu or '?'} (satır {r['_satir']})", f"Miktar '{metin(r['miktar'])}'. Satır alınmadı; müşteri listesi eksik kalabilir."))
                continue
            sevk.append({"parti": p, "parti_k": katla(p), "musteri": mu, "tarih": tarih(r["tarih"]), "miktar": m, "irsaliye": metin(r["irsaliye"])})
    return uretim, kullanim, sevk, uy


def izle(uretim: dict, kullanim: list[dict], sevk: list[dict], lot: str | None = None, parti: str | None = None) -> dict:
    if bool(lot) == bool(parti):
        raise ValueError("Hammadde lotu veya üretim partisinden yalnız biri verilmeli.")
    uy, kul_lot, kul_parti, ad = [], defaultdict(list), defaultdict(list), {}
    for k in kullanim:
        kul_lot[k["lot_k"]].append(k)
        kul_parti[k["parti_k"]].append(k)
        ad.setdefault(k["parti_k"], k["parti"])
    for s in sevk:
        ad.setdefault(s["parti_k"], s["parti"])
    for pk, u in uretim.items():
        ad[pk] = u["parti"]
    aranan = katla(lot or parti)
    etk, sira = {}, [(aranan, 0)]
    if parti:
        if aranan not in ad:
            raise ValueError(f"'{parti}' üretim, kullanım veya sevkiyat kayıtlarında bulunamadı.")
        etk[aranan] = {"seviye": 0, "kaynak": "Aranan parti"}
    elif aranan not in kul_lot:
        raise ValueError(f"'{lot}' lotu hammadde kullanım kayıtlarında bulunamadı.")
    while sira:
        a, sv = sira.pop(0)
        for k in kul_lot.get(a, []):
            if k["parti_k"] not in etk:
                etk[k["parti_k"]] = {"seviye": sv + 1, "kaynak": k["lot"]}
                sira.append((k["parti_k"], sv + 1))
    # zincir kayıtları: ileri izlemede aranan lotun ve etkilenen yarı mamullerin kullanımı; geri izlemede aranan partinin girdileri (yarı mamul üzerinden gelenler dahil)
    lotlar = []
    if parti:
        gorulen, kuyruk = set(), [(aranan, 1)]
        while kuyruk:
            a, sv = kuyruk.pop(0)
            if a in gorulen:
                continue
            gorulen.add(a)
            for k in kul_parti.get(a, []):
                yari = k["lot_k"] in uretim or k["lot_k"] in kul_parti
                lotlar.append({"yon": "Geri", "seviye": sv, "parti": k["parti"], "hammadde": k["hammadde"], "lot": k["lot"], "miktar": k["miktar"], "yari": yari})
                if yari:
                    kuyruk.append((k["lot_k"], sv + 1))
    lotlar += [{"yon": "İleri", "seviye": etk[k["parti_k"]]["seviye"], "parti": k["parti"], "hammadde": k["hammadde"], "lot": k["lot"], "miktar": k["miktar"], "yari": k["lot_k"] != aranan or bool(parti)}
               for k in kullanim if k["parti_k"] in etk and (k["lot_k"] == aranan or k["lot_k"] in etk) and etk[k["parti_k"]]["seviye"] > 0]
    partiler, sevkler, mus = [], [], {}
    for pk, e in etk.items():
        u = uretim.get(pk)
        if u is None:
            uy.append(uyari("Yüksek", "Üretim kaydı yok", ad[pk], "Parti etkilenenler arasında ama üretim dosyasında yok; üretilen ve kalan miktar hesaplanamadı."))
        sv = [s for s in sevk if s["parti_k"] == pk]
        st, ic = sum((s["miktar"] for s in sv), Z), sum((k["miktar"] or Z for k in kul_lot.get(pk, [])), Z)
        kalan = u["miktar"] - st - ic if u else None
        if kalan is not None and kalan < 0:
            uy.append(uyari("Yüksek", "Sevk üretimden fazla", ad[pk], f"Üretilen {u['miktar']:f}, sevk edilen {st:f}, iç kullanım {ic:f}: fark {kalan:f}. Parti numaraları veya miktarlar hatalı kaydedilmiş olabilir; izlenebilirlik kaydı güvenilir değil."))
        partiler.append({"parti": ad[pk], "urun": u["urun"] if u else "", "tarih": u["tarih"] if u else None, "birim": u["birim"] if u else "", "seviye": e["seviye"], "kaynak": e["kaynak"],
                         "uretilen": u["miktar"] if u else None, "sevk": st, "ic": ic, "kalan": kalan})
        for s in sv:
            sevkler.append({"parti": ad[pk], "urun": u["urun"] if u else "", "musteri": s["musteri"], "tarih": s["tarih"], "miktar": s["miktar"], "irsaliye": s["irsaliye"]})
            m = mus.setdefault(katla(s["musteri"]), {"musteri": s["musteri"], "miktar": Z, "sevkiyat": 0, "partiler": [], "ilk": None, "son": None})
            m["miktar"] += s["miktar"]
            m["sevkiyat"] += 1
            if ad[pk] not in m["partiler"]:
                m["partiler"].append(ad[pk])
            if s["tarih"]:
                m["ilk"], m["son"] = min(m["ilk"] or s["tarih"], s["tarih"]), max(m["son"] or s["tarih"], s["tarih"])
    partiler.sort(key=lambda p: (p["seviye"], p["tarih"] or date.max, p["parti"]))
    sevkler.sort(key=lambda s: (s["tarih"] or date.max, s["musteri"], s["parti"]))
    musteriler = sorted(mus.values(), key=lambda m: (-m["miktar"], m["musteri"]))
    for m in musteriler:
        m["partiler"].sort()
    if len({katla(p["birim"]) for p in partiler if p["birim"]}) > 1:
        uy.append(uyari("Bilgi", "Farklı birimler", "Özet", "Etkilenen partilerin birimleri farklı (ör. kg ve adet). Özet toplamları bu birimleri birlikte toplar; parti bazındaki satırları esas alın."))
    if not sevk:
        uy.append(uyari("Orta", "Sevkiyat verisi yok", "Sevkiyat", "Sevkiyat dosyası verilmedi veya boş; müşteri listesi çıkarılamadı."))
    ozet = {"mod": "parti" if parti else "lot", "aranan": lot or parti, "parti": len(partiler), "uretilen": sum((p["uretilen"] or Z for p in partiler), Z), "sevk": sum((p["sevk"] for p in partiler), Z),
            "ic": sum((p["ic"] for p in partiler), Z), "kalan": sum((p["kalan"] for p in partiler if p["kalan"] and p["kalan"] > 0), Z), "musteri": len(musteriler), "sevkiyat": len(sevkler)}
    return {"partiler": partiler, "musteriler": musteriler, "sevkler": sevkler, "lotlar": lotlar, "ozet": ozet, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb, o = Workbook(), r["ozet"]
    ws = wb.active
    ws.title = "Özet"
    _baslik(ws, ["Gösterge", "Değer", "Açıklama"], (32, 22, 90))
    for ad, deger, ac in (("Aranan " + ("üretim partisi" if o["mod"] == "parti" else "hammadde lotu"), o["aranan"], "İleri izleme: lotu veya partiyi içeren bütün üretim partileri"),
                          ("Etkilenen üretim partisi", o["parti"], "Yarı mamul üzerinden etkilenenler dahil"), ("Üretilen miktar", o["uretilen"], "Etkilenen partilerin toplamı"),
                          ("Sevk edilen miktar", o["sevk"], "Müşterilere giden; geri çağırma kapsamı"), ("İç kullanım", o["ic"], "Başka üretim partisinde yarı mamul olarak kullanılan"),
                          ("Depoda olması gereken", o["kalan"], "Üretilen − sevk − iç kullanım; fiili sayımla doğrulayın ve blokeye alın"),
                          ("Etkilenen müşteri", o["musteri"], ""), ("Sevkiyat sayısı", o["sevkiyat"], "")):
        n = _satir(ws, [ad, deger, ac], sar=(3,))
        if isinstance(deger, Decimal):
            ws.cell(n, 2).number_format = PF
    ws.append([])
    ws.append(["Bu rapor kayıtlardan üretilmiş bir simülasyondur. Sonuç, kayıtlarınız kadar doğrudur; eksik veya hatalı lot kaydı etkilenen ürünü gizler."])
    ws = wb.create_sheet("Etkilenen Partiler")
    _baslik(ws, ["Üretim Partisi", "Ürün", "Üretim Tarihi", "Seviye", "Bağlantı (lot / yarı mamul)", "Üretilen", "Sevk Edilen", "İç Kullanım", "Depoda Olması Gereken", "Birim", "Sayım Sonucu (siz doldurun)"],
            (16, 30, 12, 8, 22, 13, 13, 13, 14, 8, 20))
    for p in r["partiler"]:
        _satir(ws, [p["parti"], p["urun"], p["tarih"], p["seviye"], p["kaynak"], p["uretilen"], p["sevk"], p["ic"] or None, p["kalan"], p["birim"], ""], para_s=(6, 7, 8, 9), tarih_s=(3,),
               dolgu={9: KIRMIZI if p["kalan"] is None or p["kalan"] < 0 else None})
    ws = wb.create_sheet("Müşteriler")
    _baslik(ws, ["Müşteri", "Toplam Miktar", "Sevkiyat Sayısı", "Partiler", "İlk Sevk", "Son Sevk", "Bildirim Durumu (siz doldurun)"], (32, 14, 10, 40, 12, 12, 26))
    for m in r["musteriler"]:
        _satir(ws, [m["musteri"], m["miktar"], m["sevkiyat"], ", ".join(m["partiler"]), m["ilk"], m["son"], ""], para_s=(2,), tarih_s=(5, 6), sar=(4,))
    ws = wb.create_sheet("Sevkiyatlar")
    _baslik(ws, ["Sevk Tarihi", "Müşteri", "Üretim Partisi", "Ürün", "Miktar", "İrsaliye No"], (12, 32, 16, 30, 13, 16))
    for s in r["sevkler"]:
        _satir(ws, [s["tarih"], s["musteri"], s["parti"], s["urun"], s["miktar"], s["irsaliye"]], para_s=(5,), tarih_s=(1,))
    ws = wb.create_sheet("Lot Bağlantıları")
    _baslik(ws, ["Yön", "Seviye", "Üretim Partisi", "Hammadde", "Lot / Yarı Mamul Partisi", "Kullanılan Miktar", "Yarı Mamul"], (8, 8, 16, 28, 22, 14, 10))
    for k in r["lotlar"]:
        _satir(ws, [k["yon"], k["seviye"], k["parti"], k["hammadde"], k["lot"], k["miktar"], "Evet" if k["yari"] else ""], para_s=(6,))
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(uretim_yol: Path, kullanim_yol: Path, sevk_yol: Path | None, cikti: Path, lot: str | None = None, parti: str | None = None) -> dict:
    uretim, kullanim, sevk, uy = oku(uretim_yol, kullanim_yol, sevk_yol)
    r = izle(uretim, kullanim, sevk, lot, parti)
    uy += r["uyarilar"]
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Hammadde lotundan veya üretim partisinden etkilenen partileri, müşterileri ve miktarları çıkarır.")
    p.add_argument("--uretim", type=Path, default=ORNEK / "uretim.csv", help="Üretim Partisi, Ürün, Üretim Tarihi, Üretilen Miktar, Birim")
    p.add_argument("--kullanim", type=Path, default=ORNEK / "hammadde_kullanim.csv", help="Üretim Partisi, Hammadde, Hammadde Lot, Kullanılan Miktar")
    p.add_argument("--sevkiyat", type=Path, help="Üretim Partisi, Müşteri, Sevk Tarihi, Miktar, İrsaliye No")
    p.add_argument("--lot", help="İleri izleme: şüpheli hammadde lot numarası")
    p.add_argument("--parti", help="Geri ve ileri izleme: şüpheli üretim partisi numarası")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "geri_cagirma_simulasyonu.xlsx")
    a = p.parse_args(argv)
    ornek = a.uretim == ORNEK / "uretim.csv" and a.kullanim == ORNEK / "hammadde_kullanim.csv"
    sevk = a.sevkiyat or (ORNEK / "sevkiyat.csv" if ornek else None)
    lot = a.lot or ("HM-2609-A" if ornek and not a.parti else None)
    if bool(lot) == bool(a.parti):
        print("[X] --lot veya --parti seçeneklerinden yalnız biri verilmeli")
        return 2
    for y in (a.uretim, a.kullanim, sevk):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.uretim, a.kullanim, sevk, a.cikti, lot, a.parti)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {o['aranan']}: {o['parti']} üretim partisi etkilendi · üretilen {o['uretilen']:,.2f} · sevk edilen {o['sevk']:,.2f} · depoda olması gereken {o['kalan']:,.2f} · {o['musteri']} müşteri")
    for m in r["musteriler"]:
        print(f"[!] {m['musteri']}: {m['miktar']:,.2f} ({', '.join(m['partiler'])})")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
