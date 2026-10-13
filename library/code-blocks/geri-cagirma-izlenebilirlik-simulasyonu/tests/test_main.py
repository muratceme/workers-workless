import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
O = (ORNEK / "uretim.csv", ORNEK / "hammadde_kullanim.csv", ORNEK / "sevkiyat.csv")


def veri():
    return main.oku(*O)[:3]


class IzlemeTesti(unittest.TestCase):
    def test_ileri_izleme(self):
        r = main.izle(*veri(), lot="hm-2609-a")
        self.assertEqual([(p["parti"], p["seviye"], p["kaynak"], p["uretilen"], p["sevk"], p["ic"], p["kalan"]) for p in r["partiler"]],
                         [("YM-0501", 1, "HM-2609-A", D(1500), D(0), D(1400), D(100)), ("UP-1003", 1, "HM-2609-A", D(3000), D(2400), D(0), D(600)),
                          ("UP-1001", 2, "YM-0501", D(4000), D(3400), D(0), D(600)), ("UP-1004", 2, "YM-0501", D(600), D(650), D(0), D(-50))])
        self.assertEqual([(m["musteri"], m["miktar"], m["sevkiyat"], m["partiler"], m["ilk"], m["son"]) for m in r["musteriler"]],
                         [("Örnek Market A.Ş.", D(4200), 2, ["UP-1001", "UP-1003"], date(2026, 9, 20), date(2026, 9, 23)),
                          ("Numune Restoran İşletmeleri", D(1250), 2, ["UP-1003", "UP-1004"], date(2026, 9, 25), date(2026, 9, 26)),
                          ("Deneme Gıda Toptan Ltd.", D(1000), 1, ["UP-1001"], date(2026, 9, 22), date(2026, 9, 22))])
        o = r["ozet"]
        self.assertEqual((o["mod"], o["parti"], o["uretilen"], o["sevk"], o["ic"], o["kalan"], o["musteri"], o["sevkiyat"]), ("lot", 4, D(9100), D(6450), D(1400), D(1300), 3, 5))
        self.assertEqual([(k["yon"], k["seviye"], k["parti"], k["lot"], k["miktar"], k["yari"]) for k in r["lotlar"]],
                         [("İleri", 1, "YM-0501", "HM-2609-A", D(900), False), ("İleri", 2, "UP-1001", "YM-0501", D(800), True), ("İleri", 1, "UP-1003", "HM-2609-A", D(300), False), ("İleri", 2, "UP-1004", "YM-0501", D(600), True)])
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]), [("Farklı birimler", "Özet"), ("Sevk üretimden fazla", "UP-1004")])
        self.assertEqual([s["irsaliye"] for s in r["sevkler"]], ["IRS-001", "IRS-004", "IRS-006", "IRS-009", "IRS-011"])

    def test_etkilenmeyen_ve_tek_seviye(self):
        r = main.izle(*veri(), lot="HM-2610-C")
        self.assertEqual(([p["parti"] for p in r["partiler"]], [m["musteri"] for m in r["musteriler"]], r["ozet"]["kalan"], r["uyarilar"]), (["UP-1002"], ["Örnek Market A.Ş."], D(200), []))
        r = main.izle(*veri(), lot="HM-2607-B")
        self.assertEqual(([p["parti"] for p in r["partiler"]], r["ozet"]["sevk"]), (["UP-1001"], D(3400)))

    def test_geri_izleme(self):
        r = main.izle(*veri(), parti="UP-1001")
        self.assertEqual([(p["parti"], p["seviye"], p["kaynak"]) for p in r["partiler"]], [("UP-1001", 0, "Aranan parti")])
        self.assertEqual([(k["yon"], k["seviye"], k["parti"], k["lot"], k["yari"]) for k in r["lotlar"]],
                         [("Geri", 1, "UP-1001", "YM-0501", True), ("Geri", 1, "UP-1001", "HM-2607-B", False), ("Geri", 2, "YM-0501", "HM-2609-A", False), ("Geri", 2, "YM-0501", "HM-2608-T", False)])
        r = main.izle(*veri(), parti="ym-0501")
        self.assertEqual(([(p["parti"], p["seviye"]) for p in r["partiler"]], r["ozet"]["mod"], [(k["yon"], k["parti"]) for k in r["lotlar"]]),
                         ([("YM-0501", 0), ("UP-1001", 1), ("UP-1004", 1)], "parti", [("Geri", "YM-0501"), ("Geri", "YM-0501"), ("İleri", "UP-1001"), ("İleri", "UP-1004")]))

    def test_hatalar_ve_eksik_kayit(self):
        u, k, s = veri()
        for kw in ({"lot": "YOK"}, {"parti": "YOK"}, {}, {"lot": "a", "parti": "b"}):
            with self.assertRaises(ValueError):
                main.izle(u, k, s, **kw)
        del u["up 1003"]
        r = main.izle(u, k, [], lot="HM-2609-A")
        p = [p for p in r["partiler"] if p["parti"] == "UP-1003"][0]
        self.assertEqual((p["uretilen"], p["kalan"], p["urun"], sorted(x["tur"] for x in r["uyarilar"]), r["musteriler"]), (None, None, "", ["Farklı birimler", "Sevkiyat verisi yok", "Üretim kaydı yok"], []))
        k.append({"parti": "YM-0501", "parti_k": "ym 0501", "hammadde": "Rework", "lot": "UP-1001", "lot_k": "up 1001", "miktar": D(5)})   # döngü: sonsuz döngüye girmemeli
        self.assertEqual(len(main.izle(u, k, s, lot="HM-2609-A")["partiler"]), 4)
        self.assertEqual(len(main.izle(u, k, s, parti="UP-1001")["lotlar"]), 5 + 2)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "g.xlsx"
            r = main.calistir(*O, cikti, lot="HM-2609-A")
            self.assertEqual((r["ozet"]["parti"], r["ozet"]["musteri"], [u["tur"] for u in r["uyarilar"]]), (4, 3, ["Sevk üretimden fazla", "Farklı birimler"]))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Etkilenen Partiler", "Müşteriler", "Sevkiyatlar", "Lot Bağlantıları", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            u, k, s = Path(tmp) / "u.csv", Path(tmp) / "k.csv", Path(tmp) / "s.csv"
            u.write_text("Parti No;Ürün;Tarih;Miktar\nP1;A;01.10.2026;1.250,5\nP1;A;01.10.2026;49,5\nP2;B;01.10.2026;x\n", encoding="utf-8")
            k.write_text("Parti No;Bileşen;Lot;Miktar\nP1;Un;L1;12,5\nP1;Su;;1\n;;;\nP2;Un;L1;x\nP3;Un;L1;\n", encoding="utf-8")
            s.write_text("Parti No;Alıcı;Tarih;Miktar\nP1;M;02.10.2026;300\nP1;M;03.10.2026;0\n", encoding="utf-8")
            ur, ku, se, uy = main.oku(u, k, s)
            self.assertEqual((ur["p1"]["miktar"], list(ur), [(x["parti"], x["lot"], x["miktar"]) for x in ku], [(x["musteri"], x["miktar"]) for x in se], sorted(x["tur"] for x in uy)),
                             (D("1300.0"), ["p1"], [("P1", "L1", D("12.5")), ("P3", "L1", None)], [("M", D(300))],
                              ["Kullanım kaydı okunamadı", "Kullanım kaydı okunamadı", "Mükerrer üretim partisi", "Sevkiyat kaydı okunamadı", "Üretim kaydı okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--uretim", str(u), "--kullanim", str(k), "--sevkiyat", str(s), "--cikti", c]
            self.assertEqual(main.main(arg + ["--lot", "L1"]), 0)
            self.assertEqual(main.main(arg + ["--parti", "P1"]), 0)
            self.assertEqual(main.main(arg + ["--lot", "L9"]), 1)
            self.assertEqual(main.main(arg), 2)
            self.assertEqual(main.main(arg + ["--lot", "L1", "--parti", "P1"]), 2)
            self.assertEqual(main.main(["--uretim", str(Path(tmp) / "yok.csv"), "--kullanim", str(k), "--lot", "L1"]), 1)
            u.write_text("Parti No;Ürün\nP1;A\n", encoding="utf-8")
            self.assertEqual(main.main(arg + ["--lot", "L1"]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)
            self.assertEqual(main.main(["--cikti", c, "--parti", "UP-1001"]), 0)


if __name__ == "__main__":
    unittest.main()
