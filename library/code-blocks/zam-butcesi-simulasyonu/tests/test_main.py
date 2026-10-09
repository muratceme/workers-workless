import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def calistir(cikti, **kw):
    return main.calistir(ORNEK / "calisanlar.csv", ORNEK / "senaryolar.csv", cikti, ORNEK / "bantlar.csv", gecerlilik=date(2027, 1, 1), **kw)


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "z.xlsx"
        cls.s = calistir(cls.cikti, butce=D("0.25"))
        cls.sen = cls.s["senaryolar"]

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def kisi(self, senaryo, sicil):
        return next(x for x in self.sen[senaryo]["sonuclar"] if x.c.sicil == sicil)

    def test_genel_ve_seyyanen(self):
        o = self.sen["Genel %25"]["ozet"]
        self.assertEqual((o["mevcut"], o["artis"]), (D("1624900"), D("406225.00")))
        self.assertEqual(self.kisi("Seyyanen 3.000 TL + %20", "1002").yeni, D("39450.00"))   # 31000 × 1,20 + 3000 × 0,75
        self.assertNotIn("Genel %25 (bütçeye ölçekli)", self.sen)

    def test_matris(self):
        k = self.kisi("Performans matrisi", "1001")       # C, karşılaştırma oranı 0,86 → Alt → %24
        self.assertEqual((k.konum, k.oran, k.yeni), ("alt", D("0.24"), D("40300.00")))
        k = self.kisi("Performans matrisi", "1021")       # A, 59000 / 48000 = 1,23 → Üst → %26
        self.assertEqual((k.konum, k.yeni), ("ust", D("74340.00")))
        self.assertEqual(self.kisi("Performans matrisi", "1018").oran, D("0.31"))   # pozisyon bandı 82000 → Orta

    def test_isveren_maliyeti_ve_butce(self):
        o = self.sen["Genel %25"]["ozet"]
        self.assertEqual(o["yillik_isv_artis"], D("6032441.64"))
        self.assertEqual(o["butce_farki"], D("0"))
        olc = self.sen["Performans matrisi (bütçeye ölçekli)"]
        self.assertLess(abs(olc["ozet"]["butce_farki"]), 50)
        self.assertLess(olc["olcek"], 1)
        p = self.s["p"]
        self.assertEqual(main.isveren_maliyeti(D("300000"), p), D("370601.63"))   # 300000 + 297270 (SGK tavanı) × %23,75

    def test_uyarilar_ve_excel(self):
        turler = [u["tur"] for u in self.s["uyarilar"]]
        self.assertEqual(turler[0], "Asgari ücret")
        self.assertIn("Parametre", turler)
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Senaryo Karşılaştırması", "Kişi Bazında", "Departman", "Senaryo Tanımları", "Uyarılar"])
        self.assertEqual(wb["Senaryo Karşılaştırması"].max_row - 7, 5)
        self.assertEqual(wb["Kişi Bazında"]["I1"].value, "Genel %25 · Yeni Brüt")


class SecenekTesti(unittest.TestCase):
    def test_kist_bant_ustu_asgari_yuvarla(self):
        with tempfile.TemporaryDirectory() as tmp:
            s = calistir(Path(tmp) / "z.xlsx", kist=True, bant_ustu_sinirla=True, yeni_asgari=D("42000"), yuvarla=100)
            g = {x.c.sicil: x for x in s["senaryolar"]["Genel %25"]["sonuclar"]}
            self.assertEqual(g["1025"].oran, D("0.25") * 5 / 12)              # 01.08.2026 giriş → 5 ay
            self.assertEqual(g["1025"].yeni, D("130200"))                     # 117900 × 1,1042 = 130181,25 → 130200
            self.assertEqual((g["1018"].yeni, g["1018"].tek_seferlik), (D("96000.00"), D("82800.00")))   # 102900 − 96000 = 6900 × 12
            self.assertEqual((g["1001"].yeni, g["1001"].asgari_farki), (D("42000.00"), D("1300")))     # 40700 → 42000
            self.assertEqual((g["1002"].yeni, g["1002"].tek_seferlik), (D("33000.00"), D("69600.00")))   # kısmi süreli: bant üstü 44000 × 0,75

    def test_kural_yok_ve_bantsiz(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t / "c.csv").write_text("Sicil No;Brüt Ücret;Performans\n1;50000;A\n2;50000;D\n", encoding="utf-8")
            (t / "s.csv").write_text("Senaryo;Performans;Konum;Oran (%)\nM;A;Alt;30\nM;A;*;20\n", encoding="utf-8")
            s = main.calistir(t / "c.csv", t / "s.csv", t / "o.xlsx", gecerlilik=date(2026, 7, 1))
            r = s["senaryolar"]["M"]["sonuclar"]
            self.assertEqual((r[0].yeni, r[1].yeni), (D("60000.00"), D("50000")))
            turler = {u["tur"] for u in s["uyarilar"]}
            self.assertTrue({"Kural yok", "Bant yok"} <= turler)
            self.assertEqual(s["asgari"], D("33030.00"))

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx"), "--butce", "25"]), 0)
            self.assertEqual(main.main(["--senaryolar", str(Path(tmp) / "yok.csv")]), 1)
            self.assertEqual(main.main(["--gecerlilik", "2027"]), 2)


if __name__ == "__main__":
    unittest.main()
