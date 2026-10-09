"""
Zam Bütçesi Simülasyonu — Workers / Workless kod bloğu
İnsan Kaynakları › Ücretlendirme ve Yan Haklar Uzmanı

Farklı zam senaryolarında kişi bazında yeni brüt ücreti ve toplam personel maliyetini hesaplar:
  - Senaryo satırı: Performans × Konum → Oran (%) + Sabit Tutar (TL). "*" her değere uyar; en özel satır seçilir.
    Böylece genel oran ("%25"), seyyanen + oran ("3.000 TL + %20") ve performans matrisi aynı dosyada tanımlanır.
  - Konum: mevcut ücretin bant orta noktasına oranı (compa-ratio): < 0,90 Alt, 0,90–1,10 Orta, > 1,10 Üst.
  - Yeni asgari ücretin altında kalan ücret asgari ücrete çekilir (fark ayrıca gösterilir).
  - İsteğe bağlı: kıst zam (son 12 ayda girenlere çalışılan ay / 12), bant üstü sınırı (aşan kısım tek seferlik
    ödeme), yukarı yuvarlama, bütçeye ölçekleme (oran kısmı bütçeyi tutturacak şekilde orantılı küçültülür/büyütülür).
  - İşveren maliyeti: brüt + SGK işveren payı ve işsizlik işveren payı (SGK tavanı ile sınırlı; tr_parametreler.json).
Rapor: senaryo karşılaştırması, kişi bazında yeni ücretler, departman maliyeti, uygulanan matris, uyarılar.
İnternete bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: 28 çalışan, 3 senaryo
    python main.py --calisanlar calisanlar.xlsx --senaryolar senaryolar.csv --bantlar bantlar.xlsx --butce 25
    python main.py --calisanlar c.xlsx --senaryolar s.csv --yeni-asgari 40000 --kist --gecerlilik 01.01.2027
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import statistics
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import ROUND_CEILING, ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
PARAMETRE = BURASI / "tr_parametreler.json"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
K2 = Decimal("0.01")
PERF_HARF = {"5": "a", "4": "a", "3": "b", "2": "c", "1": "d"}  # 1–5 ölçeği → A/B/C/D (5-4 A, 3 B, 2 C, 1 D)

CALISAN_SUTUNLARI = {"sicil": ("sicil no", "sicil"), "ad": ("ad soyad", "ad"), "departman": ("departman", "birim"), "pozisyon": ("pozisyon", "unvan"),
                     "kademe": ("kademe", "grade", "seviye"), "ucret": ("brut ucret tl", "brut ucret", "ucret", "aylik brut"),
                     "giris": ("ise giris", "ise giris tarihi"), "performans": ("performans", "performans notu"),
                     "oran": ("calisma orani", "fte", "tam zaman orani")}
BANT_SUTUNLARI = {"kademe": ("kademe", "grade"), "pozisyon": ("pozisyon",), "alt": ("alt tl", "alt", "minimum", "p25"),
                  "orta": ("orta tl", "orta", "orta nokta", "medyan", "p50"), "ust": ("ust tl", "ust", "maksimum", "p75")}
SENARYO_SUTUNLARI = {"senaryo": ("senaryo", "senaryo adi"), "performans": ("performans", "performans notu"), "konum": ("konum", "bant konumu"),
                     "oran": ("oran", "oran", "artis orani", "zam orani"), "sabit": ("sabit tutar tl", "sabit tutar", "seyyanen", "seyyanen tl")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9*]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("%", "").replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tl(x: Decimal | None) -> str:
    return "—" if x is None else f"{x.quantize(K2, ROUND_HALF_UP):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def yuzde(x: float) -> str:
    return f"%{x * 100:.1f}".replace(".", ",")


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


def parametreler(yil: int) -> dict:
    if not PARAMETRE.exists():
        return {}
    return json.loads(PARAMETRE.read_text(encoding="utf-8")).get("yillar", {}).get(str(yil), {})


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Bant:
    kademe: str
    pozisyon: str
    alt: Decimal
    orta: Decimal
    ust: Decimal


@dataclass
class Calisan:
    sicil: str
    ad: str
    departman: str
    pozisyon: str
    kademe: str
    ucret: Decimal
    giris: date | None
    performans: str
    oran: Decimal
    bant: Bant | None = None

    @property
    def perf(self) -> str:
        k = katla(self.performans)[:1]
        return PERF_HARF.get(k, k)

    @property
    def compa(self) -> float | None:
        return float(self.ucret / self.oran / self.bant.orta) if self.bant and self.oran else None


@dataclass
class Kural:
    performans: str
    konum: str
    oran: Decimal
    sabit: Decimal

    def uyar(self, perf: str, konum: str) -> bool:
        return self.performans in ("*", "", perf) and self.konum in ("*", "", konum)

    @property
    def ozgulluk(self) -> int:
        return (self.performans not in ("*", "")) * 2 + (self.konum not in ("*", ""))


@dataclass
class Sonuc:
    c: Calisan
    konum: str
    kural: Kural | None
    oran: Decimal              # uygulanan oran (kıst sonrası, ölçek sonrası)
    yeni: Decimal
    asgari_farki: Decimal = Decimal(0)
    tek_seferlik: Decimal = Decimal(0)   # bant üstü sınırını aşan kısım, yıllık
    notlar: list = field(default_factory=list)

    @property
    def artis(self) -> Decimal:
        return self.yeni - self.c.ucret

    @property
    def artis_orani(self) -> float:
        return float(self.artis / self.c.ucret) if self.c.ucret else 0.0


def oku_calisanlar(yol: Path) -> tuple[list[Calisan], list[str]]:
    sonuc, hatalar = [], []
    for r in kayitlar(yol, CALISAN_SUTUNLARI, ("sicil", "ucret")):
        if not r.get("sicil"):
            continue
        u = para(r.get("ucret"))
        if u is None or u <= 0:
            hatalar.append(f"Satır {r['_satir']} ({r['sicil']}): ücret okunamadı")
            continue
        o = para(r.get("oran")) or Decimal(1)
        sonuc.append(Calisan(metin(r["sicil"]), metin(r.get("ad")), metin(r.get("departman")) or "—", metin(r.get("pozisyon")), metin(r.get("kademe")), u,
                             tarih(r.get("giris")), metin(r.get("performans")), o / 100 if o > 1 else o))
    return sonuc, hatalar


def oku_bantlar(yol: Path | None) -> list[Bant]:
    if not yol:
        return []
    sonuc = []
    for r in kayitlar(yol, BANT_SUTUNLARI, ("alt", "ust")):
        alt, ust = para(r.get("alt")), para(r.get("ust"))
        if alt is None or ust is None:
            continue
        sonuc.append(Bant(metin(r.get("kademe")), metin(r.get("pozisyon")), alt, para(r.get("orta")) or (alt + ust) / 2, ust))
    return sonuc


def oku_senaryolar(yol: Path) -> dict[str, list[Kural]]:
    s: dict[str, list[Kural]] = {}
    for r in kayitlar(yol, SENARYO_SUTUNLARI, ("senaryo",)):
        ad = metin(r.get("senaryo"))
        if not ad:
            continue
        oran, sabit = para(r.get("oran")), para(r.get("sabit"))
        if oran is None and sabit is None:
            raise ValueError(f"{yol.name} satır {r['_satir']}: '{ad}' için Oran (%) veya Sabit Tutar gerekli")
        s.setdefault(ad, []).append(Kural(katla(r.get("performans"))[:1] if katla(r.get("performans")) not in ("", "*") else "*",
                                          katla(r.get("konum")) or "*", (oran or Decimal(0)) / 100, sabit or Decimal(0)))
    if not s:
        raise ValueError(f"{yol.name}: senaryo bulunamadı")
    return s


# ----------------------------------------------------------------------------
# Hesap
# ----------------------------------------------------------------------------

def ay_farki(bas: date, son: date) -> int:
    """bas'tan son'a kadar tamamlanan ay sayısı."""
    a = (son.year - bas.year) * 12 + son.month - bas.month - (son.day < bas.day)
    return max(a, 0)


def konum_bul(c: Calisan, alt_esik: float, ust_esik: float) -> str:
    cr = c.compa
    if cr is None:
        return "orta"
    return "alt" if cr < alt_esik else "ust" if cr > ust_esik else "orta"


def simule_et(calisanlar: list[Calisan], kurallar: list[Kural], *, asgari: Decimal | None, gecerlilik: date, kist: bool = False,
              bant_ustu_sinirla: bool = False, bant_artis: Decimal = Decimal(0), yuvarla: int = 0, olcek: Decimal = Decimal(1),
              esikler: tuple[float, float] = (0.90, 1.10)) -> list[Sonuc]:
    sonuclar = []
    for c in calisanlar:
        konum = konum_bul(c, *esikler)
        uyan = [k for k in kurallar if k.uyar(c.perf, konum)]
        kural = max(uyan, key=lambda k: k.ozgulluk) if uyan else None
        notlar = []
        if kural is None:
            sonuclar.append(Sonuc(c, konum, None, Decimal(0), c.ucret, notlar=["Senaryoda bu performans/konum için satır yok; zam uygulanmadı"]))
            continue
        oran = kural.oran * olcek
        if kist and c.giris and c.giris > date(gecerlilik.year - 1, gecerlilik.month, min(gecerlilik.day, 28)):
            ay = ay_farki(c.giris, gecerlilik)
            oran = oran * ay / 12
            notlar.append(f"Kıst: {ay}/12 ay")
        yeni = c.ucret * (1 + oran) + kural.sabit * c.oran
        if yuvarla:
            yeni = (yeni / yuvarla).to_integral_value(ROUND_CEILING) * yuvarla
        yeni = yeni.quantize(K2, ROUND_HALF_UP)
        s = Sonuc(c, konum, kural, oran, yeni, notlar=notlar)
        if c.bant and bant_ustu_sinirla:
            ust = (c.bant.ust * (1 + bant_artis) * c.oran).quantize(K2, ROUND_HALF_UP)
            if s.yeni > ust:
                hedef = max(ust, c.ucret)
                s.tek_seferlik = (s.yeni - hedef) * 12
                notlar.append(f"Bant üstü sınırı {tl(ust)} TL; aşan {tl(s.yeni - hedef)} TL × 12 tek seferlik")
                s.yeni = hedef
        if asgari is not None and s.yeni / c.oran < asgari:
            hedef = (asgari * c.oran).quantize(K2, ROUND_HALF_UP)
            s.asgari_farki = hedef - s.yeni
            notlar.append(f"Yeni asgari ücrete çekildi (+{tl(s.asgari_farki)} TL)")
            s.yeni = hedef
        sonuclar.append(s)
    return sonuclar


def isveren_maliyeti(brut: Decimal, p: dict, tesvik: str = "yok") -> Decimal:
    if not p.get("sgk_isveren_orani"):
        return brut
    oran = Decimal(p["sgk_isveren_orani"]) + Decimal(p["issizlik_isveren_orani"]) - Decimal(p.get("sgk_isveren_tesvik_puani", {}).get(tesvik, "0"))
    matrah = min(brut, Decimal(p["sgk_tavan"])) if p.get("sgk_tavan") else brut
    return brut + (matrah * oran).quantize(K2, ROUND_HALF_UP)


def ozet(sonuclar: list[Sonuc], p: dict, tesvik: str) -> dict:
    mevcut = sum((s.c.ucret for s in sonuclar), Decimal(0))
    yeni = sum((s.yeni for s in sonuclar), Decimal(0))
    tek = sum((s.tek_seferlik for s in sonuclar), Decimal(0))
    mevcut_isv = sum((isveren_maliyeti(s.c.ucret, p, tesvik) for s in sonuclar), Decimal(0))
    yeni_isv = sum((isveren_maliyeti(s.yeni, p, tesvik) for s in sonuclar), Decimal(0))
    oranlar = [s.artis_orani for s in sonuclar]
    return {"mevcut": mevcut, "yeni": yeni, "artis": yeni - mevcut, "artis_orani": float((yeni - mevcut) / mevcut) if mevcut else 0.0,
            "ort": statistics.fmean(oranlar) if oranlar else 0.0, "medyan": statistics.median(oranlar) if oranlar else 0.0,
            "min": min(oranlar, default=0.0), "max": max(oranlar, default=0.0), "tek": tek,
            "mevcut_isv": mevcut_isv, "yeni_isv": yeni_isv, "yillik_isv_artis": (yeni_isv - mevcut_isv) * 12 + tek,
            "asgari_n": sum(1 for s in sonuclar if s.asgari_farki), "asgari_tutar": sum((s.asgari_farki for s in sonuclar), Decimal(0)),
            "bant_n": sum(1 for s in sonuclar if s.tek_seferlik)}


def butceye_olcekle(calisanlar, kurallar, butce_orani: Decimal, **kw) -> Decimal | None:
    """Toplam aylık brüt artış oranı bütçeyi tutturacak şekilde oran kısmını ölçekler (ikiye bölme)."""
    mevcut = sum((c.ucret for c in calisanlar), Decimal(0))
    hedef = mevcut * butce_orani

    def artis(k):
        return sum((s.yeni for s in simule_et(calisanlar, kurallar, olcek=k, **kw)), Decimal(0)) - mevcut

    alt, ust = Decimal(0), Decimal(5)
    if artis(alt) > hedef or artis(ust) < hedef:
        return None
    for _ in range(50):
        orta = (alt + ust) / 2
        if artis(orta) > hedef:
            ust = orta
        else:
            alt = orta
    return alt.quantize(Decimal("0.0001"))


SIRA = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}


def calistir(calisan_yolu: Path, senaryo_yolu: Path, cikti: Path, bant_yolu: Path | None = None, *, gecerlilik: date | None = None,
             yeni_asgari: Decimal | None = None, butce: Decimal | None = None, kist: bool = False, bant_ustu_sinirla: bool = False,
             bant_artis: Decimal = Decimal(0), yuvarla: int = 0, tesvik: str = "yok", param_yili: int | None = None) -> dict:
    gecerlilik = gecerlilik or date(date.today().year + 1, 1, 1)
    calisanlar, hatalar = oku_calisanlar(calisan_yolu)
    bantlar = oku_bantlar(bant_yolu)
    senaryolar = oku_senaryolar(senaryo_yolu)
    uyarilar = [{"onem": "Orta", "tur": "Okunamayan satır", "aciklama": h} for h in hatalar]
    param_yili = param_yili or gecerlilik.year
    p = parametreler(param_yili)
    if not p:
        onceki = parametreler(param_yili - 1)
        if onceki:
            uyarilar.append({"onem": "Orta", "tur": "Parametre", "aciklama": f"{param_yili} SGK oranları/tavanı tr_parametreler.json'da yok; "
                             f"{param_yili - 1} değerleri kullanıldı. Yeni dönem açıklanınca dosyayı güncelleyin."})
            p, param_yili = onceki, param_yili - 1
    asgari = yeni_asgari
    if asgari is None:
        g = parametreler(gecerlilik.year).get("asgari_ucret_brut")
        if g:
            asgari = Decimal(g)
        elif p.get("asgari_ucret_brut"):
            asgari = Decimal(p["asgari_ucret_brut"])
            uyarilar.append({"onem": "Yüksek", "tur": "Asgari ücret", "aciklama": f"{gecerlilik.year} asgari ücreti henüz tanımlı değil; taban olarak "
                             f"{param_yili} brüt asgari ücreti ({tl(asgari)} TL) kullanıldı. Açıklanınca --yeni-asgari ile yeniden çalıştırın."})
    for c in calisanlar:
        c.bant = next((b for b in bantlar if b.pozisyon and katla(b.pozisyon) == katla(c.pozisyon)), None) or \
            next((b for b in bantlar if not b.pozisyon and katla(b.kademe) == katla(c.kademe)), None)
        if bantlar and not c.bant:
            uyarilar.append({"onem": "Bilgi", "tur": "Bant yok", "aciklama": f"{c.sicil}: kademe '{c.kademe}' için bant yok; konum 'Orta' sayıldı"})
        if not c.performans and any(k.performans != "*" for ks in senaryolar.values() for k in ks):
            uyarilar.append({"onem": "Bilgi", "tur": "Performans yok", "aciklama": f"{c.sicil}: performans notu yok; senaryodaki '*' satırı uygulanır"})
    if not bantlar and any(k.konum not in ("*", "") for ks in senaryolar.values() for k in ks):
        uyarilar.append({"onem": "Orta", "tur": "Bant yok", "aciklama": "Senaryolarda Konum kullanılmış ama bant dosyası verilmedi; herkes 'Orta' sayıldı"})
    kw = dict(asgari=asgari, gecerlilik=gecerlilik, kist=kist, bant_ustu_sinirla=bant_ustu_sinirla, bant_artis=bant_artis, yuvarla=yuvarla)
    sonuc = {}
    for ad, kurallar in senaryolar.items():
        sonuc[ad] = {"kurallar": kurallar, "olcek": Decimal(1), "sonuclar": simule_et(calisanlar, kurallar, **kw)}
        if butce is not None:
            k = butceye_olcekle(calisanlar, kurallar, butce, **kw)
            if k is None:
                uyarilar.append({"onem": "Bilgi", "tur": "Bütçeye ölçekleme", "aciklama": f"'{ad}': sabit tutar ve asgari ücret farkı tek başına bütçeyi "
                                 "aşıyor veya oran 5 katına çıkarılsa da bütçeye ulaşılamıyor; ölçekli senaryo üretilmedi"})
            elif k != 1:
                sonuc[f"{ad} (bütçeye ölçekli)"] = {"kurallar": kurallar, "olcek": k, "sonuclar": simule_et(calisanlar, kurallar, olcek=k, **kw)}
    for ad, x in sonuc.items():
        x["ozet"] = ozet(x["sonuclar"], p, tesvik)
        if butce is not None:
            x["ozet"]["butce_farki"] = x["ozet"]["artis"] - x["ozet"]["mevcut"] * butce
        for s in x["sonuclar"]:
            if s.kural is None:
                uyarilar.append({"onem": "Yüksek", "tur": "Kural yok", "aciklama": f"'{ad}': {s.c.sicil} (performans {s.c.performans or '—'}, konum "
                                 f"{s.konum}) için satır yok; '*' satırı ekleyin"})
    # Senaryodan bağımsız: mevcut ücret yeni asgari ücretin altında
    for c in calisanlar:
        if asgari is not None and c.ucret / c.oran < asgari:
            uyarilar.append({"onem": "Bilgi", "tur": "Asgari ücret", "aciklama": f"{c.sicil}: mevcut ücret ({tl(c.ucret / c.oran)} TL tam zamanlı) "
                             f"taban alınan asgari ücretin ({tl(asgari)} TL) altında; senaryolarda tabana çekilir"})
    uyarilar.sort(key=lambda u: (SIRA[u["onem"]], u["tur"]))
    rapor_yaz(cikti, calisanlar, sonuc, uyarilar, {"gecerlilik": gecerlilik, "asgari": asgari, "butce": butce, "param_yili": param_yili,
                                                   "tesvik": tesvik, "kist": kist, "bant_ustu": bant_ustu_sinirla, "yuvarla": yuvarla})
    return {"calisanlar": calisanlar, "senaryolar": sonuc, "uyarilar": uyarilar, "asgari": asgari, "p": p}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KONTROL = PatternFill("solid", fgColor="FFF4CE")
UST = Alignment(vertical="top", wrap_text=True)
PF = "#,##0.00"
KONUM_AD = {"alt": "Alt", "orta": "Orta", "ust": "Üst", "*": "*", "": "*"}


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def rapor_yaz(cikti: Path, calisanlar: list[Calisan], sonuc: dict, uyarilar: list[dict], ayar: dict) -> None:
    wb = Workbook()
    ks = wb.active
    ks.title = "Senaryo Karşılaştırması"
    bas = ["Senaryo", "Ölçek", "Mevcut Aylık Brüt", "Yeni Aylık Brüt", "Aylık Artış", "Toplam Artış %", "Kişi Ort. Artış %", "Medyan %", "En Düşük %",
           "En Yüksek %", "Asgari Ücrete Çekilen", "Asgari Farkı (aylık)", "Bant Üstü Sınırlanan", "Tek Seferlik (yıllık)",
           "Yıllık İşveren Maliyeti Artışı"] + (["Bütçe Farkı (aylık)"] if ayar["butce"] is not None else []) + ["Karar"]
    _baslik(ks, bas, (30, 8, 15, 15, 14, 10, 10, 9, 9, 9, 10, 13, 10, 14, 17, 14, 18)[:len(bas)])
    for ad, x in sonuc.items():
        o = x["ozet"]
        satir = [ad, float(x["olcek"]), float(o["mevcut"]), float(o["yeni"]), float(o["artis"]), o["artis_orani"], o["ort"], o["medyan"], o["min"], o["max"],
                 o["asgari_n"], float(o["asgari_tutar"]), o["bant_n"], float(o["tek"]), float(o["yillik_isv_artis"])]
        if ayar["butce"] is not None:
            satir.append(float(o["butce_farki"]))
        ks.append(satir + [""])
        r = ks.max_row
        ks.cell(r, 2).number_format = "0.00##"
        for j in (3, 4, 5, 12, 14, 15, 16):
            ks.cell(r, j).number_format = PF
        for j in (6, 7, 8, 9, 10):
            ks.cell(r, j).number_format = "0.0%"
        if ayar["butce"] is not None:
            ks.cell(r, 16).fill = PatternFill("solid", fgColor="FDE2E1" if o["butce_farki"] > 0 else "E3F4E1")
        ks.cell(r, ks.max_column).fill = KONTROL
    ks.append([])
    bilgi = [f"Geçerlilik tarihi: {ayar['gecerlilik']:%d.%m.%Y}",
             f"Taban (asgari ücret, brüt): {tl(ayar['asgari'])} TL" if ayar["asgari"] else "Taban: asgari ücret uygulanmadı",
             f"İşveren maliyeti: {ayar['param_yili']} SGK işveren + işsizlik işveren payı, SGK tavanı ile sınırlı; teşvik: {ayar['tesvik']}. "
             "Yıllık = aylık × 12 + tek seferlik; ikramiye, yan haklar ve kıdem karşılığı dahil değildir.",
             f"Bütçe: toplam aylık brütün %{ayar['butce'] * 100:g}'i".replace(".", ",") if ayar["butce"] is not None else "Bütçe verilmedi (--butce)",
             "Seçenekler: " + ", ".join(x for x in ["kıst zam" if ayar["kist"] else "", "bant üstü sınırı" if ayar["bant_ustu"] else "",
                                                     f"{ayar['yuvarla']} TL'ye yukarı yuvarlama" if ayar["yuvarla"] else ""] if x) or "Seçenekler: yok"]
    for b in bilgi:
        ks.append([b])

    kb = wb.create_sheet("Kişi Bazında")
    adlar = list(sonuc)
    bas = ["Sicil", "Ad Soyad", "Departman", "Kademe", "Performans", "Karşılaştırma Oranı", "Konum", "Mevcut Brüt"]
    for ad in adlar:
        bas += [f"{ad} · Yeni Brüt", f"{ad} · Artış %", f"{ad} · Not"]
    _baslik(kb, bas + ["Karar"], (8, 16, 14, 7, 9, 10, 7, 12) + (13, 9, 26) * len(adlar) + (16,))
    ilk = sonuc[adlar[0]]["sonuclar"]
    for i, c in enumerate(calisanlar):
        satir = [c.sicil, c.ad, c.departman, c.kademe, c.performans, c.compa, KONUM_AD[ilk[i].konum], float(c.ucret)]
        for ad in adlar:
            s = sonuc[ad]["sonuclar"][i]
            satir += [float(s.yeni), s.artis_orani, "; ".join(s.notlar)]
        kb.append(satir + [""])
        r = kb.max_row
        kb.cell(r, 6).number_format = "0.00"
        kb.cell(r, 8).number_format = PF
        for k in range(len(adlar)):
            kb.cell(r, 9 + 3 * k).number_format = PF
            kb.cell(r, 10 + 3 * k).number_format = "0.0%"
            if sonuc[adlar[k]]["sonuclar"][i].notlar:
                kb.cell(r, 11 + 3 * k).fill = PatternFill("solid", fgColor="FFF4CE")
        kb.cell(r, kb.max_column).fill = KONTROL
    kb.auto_filter.ref = f"A1:{get_column_letter(kb.max_column)}{kb.max_row}"

    dp = wb.create_sheet("Departman")
    _baslik(dp, ["Departman", "Çalışan", "Mevcut Aylık Brüt"] + [f"{ad} · Aylık Artış" for ad in adlar] + [f"{ad} · Artış %" for ad in adlar],
            (18, 8, 15) + (16,) * len(adlar) * 2)
    g = defaultdict(list)
    for i, c in enumerate(calisanlar):
        g[c.departman].append(i)
    for dep, idx in sorted(g.items()):
        mevcut = sum((calisanlar[i].ucret for i in idx), Decimal(0))
        artislar = [sum((sonuc[ad]["sonuclar"][i].artis for i in idx), Decimal(0)) for ad in adlar]
        dp.append([dep, len(idx), float(mevcut)] + [float(a) for a in artislar] + [float(a / mevcut) for a in artislar])
        r = dp.max_row
        for j in range(3, 4 + len(adlar)):
            dp.cell(r, j).number_format = PF
        for j in range(4 + len(adlar), 4 + 2 * len(adlar)):
            dp.cell(r, j).number_format = "0.0%"

    mt = wb.create_sheet("Senaryo Tanımları")
    _baslik(mt, ["Senaryo", "Performans", "Konum", "Oran (tanım)", "Ölçek", "Uygulanan Oran", "Sabit Tutar (TL)", "Uyan Çalışan"], (30, 11, 9, 11, 8, 13, 14, 11))
    for ad, x in sonuc.items():
        for k in x["kurallar"]:
            n = sum(1 for s in x["sonuclar"] if s.kural is k)
            mt.append([ad, k.performans.upper() if k.performans != "*" else "*", KONUM_AD.get(k.konum, k.konum), float(k.oran), float(x["olcek"]),
                       float(k.oran * x["olcek"]), float(k.sabit), n])
            for j in (4, 6):
                mt.cell(mt.max_row, j).number_format = "0.00%"
            mt.cell(mt.max_row, 7).number_format = PF
    mt.append([])
    mt.append(["Konum: karşılaştırma oranı (mevcut ücret ÷ bant orta noktası) < 0,90 Alt, 0,90–1,10 Orta, > 1,10 Üst. Performans 1–5 ise 5-4 A, 3 B, 2 C, 1 D."])

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Açıklama"], (9, 18, 110))
    for u in uyarilar:
        uy.append([u["onem"], u["tur"], u["aciklama"]])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 3).alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Zam senaryolarında kişi bazında yeni ücretleri ve toplam personel maliyetini hesaplar.")
    p.add_argument("--calisanlar", type=Path, default=ORNEK / "calisanlar.csv", help="Sicil, Ad, Departman, Kademe, Brüt Ücret, İşe Giriş, Performans, …")
    p.add_argument("--senaryolar", type=Path, default=ORNEK / "senaryolar.csv", help="Senaryo, Performans, Konum, Oran (%%), Sabit Tutar (TL)")
    p.add_argument("--bantlar", type=Path, help="İsteğe bağlı ücret bantları (Kademe, Pozisyon, Alt, Orta, Üst); konum için gerekir")
    p.add_argument("--gecerlilik", help="Zammın geçerlilik tarihi GG.AA.YYYY (varsayılan gelecek yılın 1 Ocak'ı)")
    p.add_argument("--yeni-asgari", help="Geçerlilik dönemindeki brüt asgari ücret (açıklanmadıysa önceki yıl kullanılır)")
    p.add_argument("--butce", type=float, help="Aylık brüt artış bütçesi, toplam brütün %%'si (ör. 25); bütçeye ölçekli senaryolar da üretilir")
    p.add_argument("--kist", action="store_true", help="Son 12 ayda girenlere oran kısmını çalışılan ay / 12 oranında uygula")
    p.add_argument("--bant-ustu-sinirla", action="store_true", help="Yeni ücret bant üstünü aşarsa üstte sınırla, aşan kısmı tek seferlik say")
    p.add_argument("--bant-artis", type=float, default=0, help="Bant üst sınırının yeni dönemde beklenen artışı %% (varsayılan 0)")
    p.add_argument("--yuvarla", type=int, default=0, help="Yeni ücreti bu tutarın katına yukarı yuvarla (ör. 100)")
    p.add_argument("--tesvik", choices=("yok", "imalat", "diger"), default="yok", help="SGK işveren payı teşvik puanı (tr_parametreler.json)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "zam_butcesi_simulasyonu.xlsx")
    a = p.parse_args(argv)
    for y in (a.calisanlar, a.senaryolar, a.bantlar):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    ornek = a.calisanlar == ORNEK / "calisanlar.csv"
    bant = a.bantlar or (ORNEK / "bantlar.csv" if ornek else None)
    gec = tarih(a.gecerlilik) if a.gecerlilik else (date(2027, 1, 1) if ornek else None)
    if a.gecerlilik and not gec:
        print("[X] --gecerlilik GG.AA.YYYY biçiminde olmalı")
        return 2
    try:
        s = calistir(a.calisanlar, a.senaryolar, a.cikti, bant, gecerlilik=gec, yeni_asgari=para(a.yeni_asgari) if a.yeni_asgari else None,
                     butce=Decimal(str(a.butce)) / 100 if a.butce is not None else None, kist=a.kist, bant_ustu_sinirla=a.bant_ustu_sinirla,
                     bant_artis=Decimal(str(a.bant_artis)) / 100, yuvarla=a.yuvarla, tesvik=a.tesvik)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(s['calisanlar'])} çalışan · {len(s['senaryolar'])} senaryo")
    for ad, x in s["senaryolar"].items():
        o = x["ozet"]
        print(f"     {ad}: aylık artış {tl(o['artis'])} TL ({yuzde(o['artis_orani'])}), yıllık işveren maliyeti artışı {tl(o['yillik_isv_artis'])} TL"
              + (f", bütçe farkı {tl(o['butce_farki'])} TL/ay" if "butce_farki" in o else ""))
    for u in s["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
