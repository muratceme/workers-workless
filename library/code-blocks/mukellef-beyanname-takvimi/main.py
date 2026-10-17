"""
Mükellef Beyanname Takvimi — Gerisi Bende kod bloğu
Beyanname ve Vergi › Serbest Muhasebeci Mali Müşavir (SMMM)

Mükelleflerin yükümlülüklerinden seçilen ayın beyanname ve ödeme takvimini çıkarır; verilen ve bekleyenleri izler.
  Beyanname türü     = sizin dosyanızdan: periyot (aylık, 3 aylık, yıllık), dönem sonu ayları, dönemden kaç ay sonra verildiği, beyan ve ödeme günü
  Bu ayın işleri     = takvim ayından "kaç ay sonra" kadar geriye gidilen ay, türün dönem sonu aylarından biriyse o beyanname bu ay verilir
  Son gün            = ayın beyan günü ("son" ise ayın son günü); cumartesi, pazar veya tatil dosyanızdaki bir güne rastlarsa izleyen ilk iş günü
  Durum              = Verildi · Geç verildi · Gecikmiş (son gün geçti, kayıt yok) · Yaklaşıyor (--yaklasan gün veya daha az) · Bekliyor
  Ödeme durumu       = Ödendi · Geç ödendi · Gecikmiş · Yaklaşıyor · Bekliyor                    (ödeme günü tanımlıysa)
Kod vergi takvimini bilmez: günleri siz girersiniz. Örnek dosyadaki günler örnektir; süreler tebliğle değişir, uzatılır. Her ay GİB vergi takviminden doğrulayın.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
import calendar

TUR = {"ad": ("beyanname", "beyanname turu", "yukumluluk", "bildirim"), "periyot": ("periyot", "donem tipi", "siklik"), "aylar": ("donem sonu aylari", "donem aylari", "aylar"),
       "sonra": ("kac ay sonra", "ay sonra", "gecikme ayi"), "beyan": ("beyan gunu", "son gun", "beyanname gunu"), "odeme": ("odeme gunu", "odeme")}
YUK = {"mukellef": ("mukellef", "musteri", "firma", "unvan"), "ad": TUR["ad"], "sorumlu": ("sorumlu", "sorumlu personel")}
DURUM = {"mukellef": YUK["mukellef"], "ad": TUR["ad"], "donem": ("donem", "ait oldugu donem"), "verildi": ("verildi", "verilis tarihi", "beyan tarihi", "tahakkuk tarihi"), "odendi": ("odendi", "odeme tarihi")}
TATIL = {"tarih": ("tarih", "tatil", "gun"), "ad": ("aciklama", "tatil adi", "ad")}
AYLAR = ("Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık")
AY_NO = {katla(a): i for i, a in enumerate(AYLAR, 1)}
PERIYOT = {"aylik": set(range(1, 13)), "3 aylik": {3, 6, 9, 12}, "uc aylik": {3, 6, 9, 12}, "ceyrek": {3, 6, 9, 12}, "yillik": {12}}
RENK_D = {"Gecikmiş": KIRMIZI, "Yaklaşıyor": SARI, "Verildi": YESIL, "Ödendi": YESIL, "Geç verildi": MAVI, "Geç ödendi": MAVI}


def donem_coz(x) -> tuple[int, int] | None:
    """2026-09 · 09.2026 · 09/2026 · Eylül 2026 · tarih hücresi → (yıl, ay)"""
    if isinstance(x, (date, datetime)):
        return x.year, x.month
    s = metin(x)
    m = re.fullmatch(r"(\d{4})[-./](\d{1,2})(?:[-./]\d{1,2})?", s)
    if m:
        y, a = int(m.group(1)), int(m.group(2))
    else:
        m = re.fullmatch(r"(?:\d{1,2}[-./])?(\d{1,2})[-./](\d{4})", s)
        if m:
            y, a = int(m.group(2)), int(m.group(1))
        else:
            p = katla(s).split()
            if len(p) != 2 or p[0] not in AY_NO or not p[1].isdigit():
                return None
            y, a = int(p[1]), AY_NO[p[0]]
    return (y, a) if 1 <= a <= 12 and 2000 <= y <= 2100 else None


def gun_coz(x) -> int | str | None:
    s = katla(x)
    if s in ("son", "son gun", "ay sonu"):
        return "son"
    return int(s) if s.isdigit() and 1 <= int(s) <= 31 else None


def oku(tur_yol: Path, yuk_yol: Path, durum_yol: Path | None = None, tatil_yol: Path | None = None) -> tuple[dict, list[dict], dict, set, list[dict]]:
    uy, turler, yukler, durum, tatiller, gorulen = [], {}, [], {}, set(), set()
    for r in kayitlar(tur_yol, TUR, ("ad", "periyot", "beyan")):
        ad = metin(r["ad"])
        if not ad:
            continue
        aylar = PERIYOT.get(katla(r["periyot"]))
        if metin(r["aylar"]):
            ozel = re.findall(r"\d+", metin(r["aylar"]))
            aylar = {int(a) for a in ozel} if ozel and all(1 <= int(a) <= 12 for a in ozel) else None
        sonra, beyan, odeme = (int(metin(r["sonra"])) if metin(r["sonra"]).isdigit() else None if metin(r["sonra"]) else 1), gun_coz(r["beyan"]), gun_coz(r["odeme"]) if metin(r["odeme"]) else None
        if not aylar or sonra is None or not 0 <= sonra <= 12 or beyan is None or (metin(r["odeme"]) and odeme is None):
            uy.append(uyari("Yüksek", "Beyanname türü okunamadı", ad, f"Periyot '{metin(r['periyot'])}', dönem sonu ayları '{metin(r['aylar'])}', kaç ay sonra '{metin(r['sonra'])}', beyan günü '{metin(r['beyan'])}', ödeme günü '{metin(r['odeme'])}'. "
                            "Periyot Aylık, 3 Aylık veya Yıllık; günler 1–31 veya 'son' olmalı. Bu türün yükümlülükleri takvime alınmadı."))
            continue
        turler[katla(ad)] = {"ad": ad, "periyot": metin(r["periyot"]), "aylar": aylar, "sonra": sonra, "beyan": beyan, "odeme": odeme}
    for r in kayitlar(yuk_yol, YUK, ("mukellef", "ad")):
        m, ad = metin(r["mukellef"]), metin(r["ad"])
        if not m or not ad:
            continue
        if (katla(m), katla(ad)) in gorulen:
            continue
        gorulen.add((katla(m), katla(ad)))
        if katla(ad) not in turler:
            uy.append(uyari("Yüksek", "Tanımsız beyanname türü", f"{m} · {ad}", "Bu beyanname türü türler dosyasında yok veya okunamadı; mükellefin bu yükümlülüğü takvime alınmadı."))
            continue
        yukler.append({"mukellef": m, "tur": katla(ad), "sorumlu": metin(r["sorumlu"])})
    if durum_yol:
        for r in kayitlar(durum_yol, DURUM, ("mukellef", "ad", "donem")):
            m, ad = metin(r["mukellef"]), metin(r["ad"])
            if not m and not ad:
                continue
            d, v, o = donem_coz(r["donem"]), tarih(r["verildi"]), tarih(r["odendi"])
            if d is None or (v is None and metin(r["verildi"])) or (o is None and metin(r["odendi"])):
                uy.append(uyari("Yüksek", "Durum satırı okunamadı", f"{m} · {ad} (satır {r['_satir']})", f"Dönem '{metin(r['donem'])}', verildi '{metin(r['verildi'])}', ödendi '{metin(r['odendi'])}'. Satır alınmadı; beyanname verilmemiş görünebilir."))
                continue
            durum[(katla(m), katla(ad), d)] = {"verildi": v, "odendi": o}
    if tatil_yol:
        for r in kayitlar(tatil_yol, TATIL, ("tarih",)):
            t = tarih(r["tarih"])
            if t:
                tatiller.add(t)
    return turler, yukler, durum, tatiller, uy


def is_gunu(t: date, tatiller: set) -> date:
    while t.weekday() >= 5 or t in tatiller:
        t += timedelta(days=1)
    return t


def son_gun(yil: int, ay: int, gun: int | str, tatiller: set) -> date:
    son = calendar.monthrange(yil, ay)[1]
    return is_gunu(date(yil, ay, son if gun == "son" else min(gun, son)), tatiller)


def hal(son: date, yapildi: date | None, bugun: date, yaklasan: int, tamam: str) -> str:
    if yapildi:
        return tamam if yapildi <= son else {"Verildi": "Geç verildi", "Ödendi": "Geç ödendi"}[tamam]
    return "Gecikmiş" if bugun > son else "Yaklaşıyor" if (son - bugun).days <= yaklasan else "Bekliyor"


def takvim(turler: dict, yukler: list[dict], ay: tuple[int, int], bugun: date, durum: dict | None = None, tatiller: set | None = None, yaklasan: int = 3) -> dict:
    if not yukler:
        raise ValueError("Takvime alınacak yükümlülük yok.")
    durum, tatiller, uy, (y, m) = durum or {}, tatiller or set(), [], ay
    satirlar = []
    for yk in yukler:
        t = turler[yk["tur"]]
        dm = (y * 12 + m - 1) - t["sonra"]
        dy, da = divmod(dm, 12)
        da += 1
        if da not in t["aylar"]:
            continue
        donem_adi = f"{AYLAR[da - 1]} {dy}" if len(t["aylar"]) == 12 else f"{dy}" if t["aylar"] == {12} else f"{dy} {AYLAR[da - 1]} dönemi"
        d = durum.get((katla(yk["mukellef"]), yk["tur"], (dy, da)), {"verildi": None, "odendi": None})
        beyan = son_gun(y, m, t["beyan"], tatiller)
        odeme = son_gun(y, m, t["odeme"], tatiller) if t["odeme"] else None
        x = {"mukellef": yk["mukellef"], "beyanname": t["ad"], "donem": donem_adi, "sorumlu": yk["sorumlu"], "beyan": beyan, "odeme": odeme, "verildi": d["verildi"], "odendi": d["odendi"],
             "durum": hal(beyan, d["verildi"], bugun, yaklasan, "Verildi"), "odeme_durum": hal(odeme, d["odendi"], bugun, yaklasan, "Ödendi") if odeme else "", "kalan": (beyan - bugun).days}
        satirlar.append(x)
        if x["durum"] == "Gecikmiş":
            uy.append(uyari("Yüksek", "Beyanname gecikmiş", f"{x['mukellef']} · {t['ad']}", f"{donem_adi} dönemi için son gün {beyan:%d.%m.%Y} idi; verildi kaydı yok ({-x['kalan']} gün geçti)."))
        elif x["durum"] == "Yaklaşıyor":
            uy.append(uyari("Orta", "Beyanname yaklaşıyor", f"{x['mukellef']} · {t['ad']}", f"{donem_adi} dönemi için son gün {beyan:%d.%m.%Y} ({x['kalan']} gün kaldı)."))
        if x["odeme_durum"] == "Gecikmiş":
            uy.append(uyari("Yüksek", "Ödeme gecikmiş", f"{x['mukellef']} · {t['ad']}", f"{donem_adi} dönemi ödemesinin son günü {odeme:%d.%m.%Y} idi; ödendi kaydı yok."))
    if not satirlar:
        raise ValueError(f"{AYLAR[m - 1]} {y} ayında verilecek beyanname bulunamadı.")
    satirlar.sort(key=lambda x: (x["beyan"], x["beyanname"], x["mukellef"]))
    ozet = {}
    for x in satirlar:
        o = ozet.setdefault((x["beyan"], x["beyanname"]), {"tarih": x["beyan"], "gun": ("Pzt", "Sal", "Çar", "Per", "Cum", "Cmt", "Paz")[x["beyan"].weekday()], "beyanname": x["beyanname"], "donem": x["donem"], "odeme": x["odeme"],
                                                           "mukellef": 0, "verilen": 0, "kalan": 0, "geciken": 0})
        o["mukellef"] += 1
        o["verilen"] += x["durum"] in ("Verildi", "Geç verildi")
        o["kalan"] += x["durum"] in ("Bekliyor", "Yaklaşıyor", "Gecikmiş")
        o["geciken"] += x["durum"] == "Gecikmiş"
    say = defaultdict(int)
    for x in satirlar:
        say[x["durum"]] += 1
    return {"satirlar": satirlar, "ozet": list(ozet.values()), "say": dict(say), "ay": f"{AYLAR[m - 1]} {y}", "bugun": bugun, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Takvim"
    _baslik(ws, ["Son Gün", "Gün", "Beyanname", "Dönem", "Ödeme Son Günü", "Mükellef Sayısı", "Verilen", "Kalan", "Geciken"], (12, 6, 34, 20, 14, 10, 9, 9, 9))
    for o in r["ozet"]:
        _satir(ws, [o["tarih"], o["gun"], o["beyanname"], o["donem"], o["odeme"], o["mukellef"], o["verilen"] or None, o["kalan"] or None, o["geciken"] or None], tarih_s=(1, 5),
               dolgu={9: KIRMIZI if o["geciken"] else None, 8: YESIL if not o["kalan"] else None})
    ws.append([])
    ws.append([f"{r['ay']} takvimi · rapor tarihi {r['bugun']:%d.%m.%Y}. Günler türler dosyanızdan gelir; hafta sonu ve tatil dosyanızdaki günler izleyen iş gününe kaydırılır. Süre uzatımlarını ve güncel vergi takvimini GİB'den doğrulayın."])
    ws = wb.create_sheet("Mükellef Bazında")
    _baslik(ws, ["Mükellef", "Beyanname", "Dönem", "Son Gün", "Kalan Gün", "Durum", "Verildiği Tarih", "Ödeme Son Günü", "Ödeme Durumu", "Ödendiği Tarih", "Sorumlu", "Not (siz doldurun)"], (32, 34, 20, 12, 9, 13, 13, 14, 13, 13, 18, 28))
    for x in sorted(r["satirlar"], key=lambda x: (x["mukellef"], x["beyan"], x["beyanname"])):
        _satir(ws, [x["mukellef"], x["beyanname"], x["donem"], x["beyan"], x["kalan"], x["durum"], x["verildi"], x["odeme"], x["odeme_durum"], x["odendi"], x["sorumlu"], ""], tarih_s=(4, 7, 8, 10),
               dolgu={6: RENK_D.get(x["durum"]), 9: RENK_D.get(x["odeme_durum"])})
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(tur_yol: Path, yuk_yol: Path, cikti: Path, ay: tuple[int, int], bugun: date, durum_yol: Path | None = None, tatil_yol: Path | None = None, yaklasan: int = 3) -> dict:
    turler, yukler, durum, tatiller, uy = oku(tur_yol, yuk_yol, durum_yol, tatil_yol)
    r = takvim(turler, yukler, ay, bugun, durum, tatiller, yaklasan)
    uy += r["uyarilar"]
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Mükellef yükümlülüklerinden seçilen ayın beyanname ve ödeme takvimini çıkarır.")
    p.add_argument("--yukumlulukler", type=Path, default=ORNEK / "yukumlulukler.csv", help="Mükellef, Beyanname, Sorumlu")
    p.add_argument("--turler", type=Path, help="Beyanname, Periyot, Dönem Sonu Ayları, Kaç Ay Sonra, Beyan Günü, Ödeme Günü")
    p.add_argument("--durum", type=Path, help="İsteğe bağlı: Mükellef, Beyanname, Dönem, Verildi, Ödendi")
    p.add_argument("--tatiller", type=Path, help="İsteğe bağlı: Tarih, Açıklama (resmi tatil günleri)")
    p.add_argument("--ay", help="Takvim ayı (ör. 2026-10); verilmezse bugünün ayı")
    p.add_argument("--tarih", help="Bugünün tarihi (GG.AA.YYYY); verilmezse sistem tarihi")
    p.add_argument("--yaklasan", type=int, default=3, help="Son güne bu kadar gün veya daha az kalanlar 'Yaklaşıyor' (varsayılan 3)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "beyanname_takvimi.xlsx")
    a = p.parse_args(argv)
    ornek = a.yukumlulukler == ORNEK / "yukumlulukler.csv"
    tur, durum, tatil = a.turler or (ORNEK / "beyanname_turleri.csv" if ornek else None), a.durum or (ORNEK / "durum.csv" if ornek else None), a.tatiller or (ORNEK / "tatiller.csv" if ornek else None)
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 26) if ornek and not a.ay else date.today()
    ay = donem_coz(a.ay) if a.ay else (bugun.year, bugun.month) if bugun else None
    if tur is None or bugun is None or ay is None or a.yaklasan < 0:
        print("[X] --turler dosyası verilmeli; --ay 2026-10 gibi, --tarih GG.AA.YYYY biçiminde, --yaklasan sıfır veya pozitif olmalı")
        return 2
    for y in (a.yukumlulukler, tur, durum, tatil):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(tur, a.yukumlulukler, a.cikti, ay, bugun, durum, tatil, a.yaklasan)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    s = r["say"]
    print(f"[OK] {r['ay']} · {len(r['satirlar'])} beyanname · " + " · ".join(f"{n} {k.lower()}" for k, n in sorted(s.items())))
    for o in r["ozet"]:
        print(f"[{'!' if o['geciken'] else 'OK'}] {o['tarih']:%d.%m.%Y} {o['gun']} · {o['beyanname']} ({o['donem']}): {o['mukellef']} mükellef, {o['kalan']} kalan")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
