import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
B = date(2026, 10, 26)


def tur(ad, aylar, sonra, beyan, odeme=None):
    return {"ad": ad, "periyot": "", "aylar": set(aylar), "sonra": sonra, "beyan": beyan, "odeme": odeme}


TUM, CEYREK = range(1, 13), (3, 6, 9, 12)
TURLER = {"kdv": tur("KDV", TUM, 1, 28, 28), "muh": tur("Muhtasar", TUM, 1, 26, 26), "gv": tur("Geçici", CEYREK, 2, 17, 17), "babs": tur("Ba-Bs", TUM, 1, "son"), "kv": tur("Kurumlar", (12,), 4, 30, 30)}
yk = lambda m, t: {"mukellef": m, "tur": t, "sorumlu": ""}   # noqa: E731
YUK = [yk("A", t) for t in TURLER] + [yk("B", "kdv"), yk("B", "muh")]


class TakvimTesti(unittest.TestCase):
    def test_ayin_isleri_ve_son_gunler(self):
        r = main.takvim(TURLER, YUK, (2026, 10), B)
        self.assertEqual([(x["mukellef"], x["beyanname"], x["donem"], x["beyan"], x["odeme"], x["kalan"], x["durum"], x["odeme_durum"]) for x in r["satirlar"]],
                         [("A", "Muhtasar", "Eylül 2026", B, B, 0, "Yaklaşıyor", "Yaklaşıyor"), ("B", "Muhtasar", "Eylül 2026", B, B, 0, "Yaklaşıyor", "Yaklaşıyor"),
                          ("A", "KDV", "Eylül 2026", date(2026, 10, 28), date(2026, 10, 28), 2, "Yaklaşıyor", "Yaklaşıyor"), ("B", "KDV", "Eylül 2026", date(2026, 10, 28), date(2026, 10, 28), 2, "Yaklaşıyor", "Yaklaşıyor"),
                          ("A", "Ba-Bs", "Eylül 2026", date(2026, 11, 2), None, 7, "Bekliyor", "")])                    # 31.10.2026 cumartesi → pazartesi
        self.assertEqual([(o["tarih"].day, o["gun"], o["beyanname"], o["mukellef"], o["verilen"], o["kalan"]) for o in r["ozet"]], [(26, "Pzt", "Muhtasar", 2, 0, 2), (28, "Çar", "KDV", 2, 0, 2), (2, "Pzt", "Ba-Bs", 1, 0, 1)])
        r = main.takvim(TURLER, YUK, (2026, 11), B)
        self.assertEqual([(x["beyanname"], x["donem"], x["beyan"]) for x in r["satirlar"] if x["mukellef"] == "A"],
                         [("Geçici", "2026 Eylül dönemi", date(2026, 11, 17)), ("Muhtasar", "Ekim 2026", date(2026, 11, 26)), ("Ba-Bs", "Ekim 2026", date(2026, 11, 30)), ("KDV", "Ekim 2026", date(2026, 11, 30))])   # 28.11 cumartesi
        r = main.takvim(TURLER, YUK, (2027, 4), B)
        self.assertEqual([(x["beyanname"], x["donem"], x["beyan"]) for x in r["satirlar"] if x["beyanname"] in ("Kurumlar", "KDV") and x["mukellef"] == "A"],
                         [("KDV", "Mart 2027", date(2027, 4, 28)), ("Kurumlar", "2026", date(2027, 4, 30))])
        r = main.takvim(TURLER, YUK, (2027, 2), B)
        self.assertEqual([(x["beyanname"], x["donem"], x["beyan"]) for x in r["satirlar"] if x["mukellef"] == "A"],
                         [("Geçici", "2026 Aralık dönemi", date(2027, 2, 17)), ("Muhtasar", "Ocak 2027", date(2027, 2, 26)), ("Ba-Bs", "Ocak 2027", date(2027, 3, 1)), ("KDV", "Ocak 2027", date(2027, 3, 1))])   # 28.02.2027 pazar

    def test_durum_tatil_ve_uyarilar(self):
        durum = {("a", "kdv", (2026, 9)): {"verildi": date(2026, 10, 24), "odendi": None}, ("a", "muh", (2026, 9)): {"verildi": date(2026, 10, 27), "odendi": date(2026, 10, 26)},
                 ("b", "kdv", (2026, 8)): {"verildi": date(2026, 9, 20), "odendi": None}}
        r = main.takvim(TURLER, YUK, (2026, 10), date(2026, 10, 29), durum, {date(2026, 10, 28), date(2026, 10, 29)}, yaklasan=5)
        d = {(x["mukellef"], x["beyanname"]): (x["beyan"].day, x["durum"], x["odeme_durum"]) for x in r["satirlar"]}
        self.assertEqual(d, {("A", "Muhtasar"): (26, "Geç verildi", "Ödendi"), ("B", "Muhtasar"): (26, "Gecikmiş", "Gecikmiş"), ("A", "KDV"): (30, "Verildi", "Yaklaşıyor"), ("B", "KDV"): (30, "Yaklaşıyor", "Yaklaşıyor"),
                             ("A", "Ba-Bs"): (2, "Yaklaşıyor", "")})
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]), [("Beyanname gecikmiş", "B · Muhtasar"), ("Beyanname yaklaşıyor", "A · Ba-Bs"), ("Beyanname yaklaşıyor", "B · KDV"), ("Ödeme gecikmiş", "B · Muhtasar")])
        self.assertEqual(r["say"], {"Geç verildi": 1, "Gecikmiş": 1, "Verildi": 1, "Yaklaşıyor": 2})
        for yukler, ay in (([], (2026, 10)), ([yk("A", "gv")], (2026, 10))):
            with self.assertRaises(ValueError):
                main.takvim(TURLER, yukler, ay, B)

    def test_yardimcilar(self):
        self.assertEqual([main.donem_coz(x) for x in ("2026-09", "09.2026", "Eylül 2026", "2026-09-15", "x", "2026-13")], [(2026, 9)] * 4 + [None, None])
        self.assertEqual([main.gun_coz(x) for x in ("28", "son", "Ay Sonu", "0", "32", "x", 26)], [28, "son", "son", None, None, None, 26])
        self.assertEqual((main.son_gun(2026, 2, 30, set()), main.son_gun(2026, 5, "son", set()), main.son_gun(2026, 10, 31, {date(2026, 11, 2)})), (date(2026, 3, 2), date(2026, 6, 1), date(2026, 11, 3)))


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "b.xlsx"
            r = main.calistir(ORNEK / "beyanname_turleri.csv", ORNEK / "yukumlulukler.csv", cikti, (2026, 10), B, ORNEK / "durum.csv", ORNEK / "tatiller.csv")
            self.assertEqual([(o["tarih"], o["beyanname"], o["mukellef"], o["verilen"], o["kalan"]) for o in r["ozet"]],
                             [(B, "Muhtasar ve Prim Hizmet Beyannamesi", 2, 1, 1), (date(2026, 10, 28), "KDV Beyannamesi", 3, 1, 2), (date(2026, 11, 2), "Form Ba-Bs", 1, 0, 1)])
            self.assertEqual((r["say"], [(u["onem"], u["tur"]) for u in r["uyarilar"]]), ({"Verildi": 2, "Yaklaşıyor": 3, "Bekliyor": 1}, [("Yüksek", "Tanımsız beyanname türü")] + [("Orta", "Beyanname yaklaşıyor")] * 3))
            x = [s for s in r["satirlar"] if s["mukellef"].startswith("Örnek") and s["beyanname"] == "KDV Beyannamesi"][0]
            self.assertEqual((x["durum"], x["odeme_durum"], x["verildi"]), ("Verildi", "Yaklaşıyor", date(2026, 10, 24)))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Takvim", "Mükellef Bazında", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            ty, yy, dy = Path(tmp) / "t.csv", Path(tmp) / "y.csv", Path(tmp) / "d.csv"
            ty.write_text("Beyanname;Periyot;Dönem Sonu Ayları;Kaç Ay Sonra;Beyan Günü;Ödeme Günü\nKDV;Aylık;;;28;28\nÖzel;Yıllık;6;2;son;\nX;Haftalık;;1;5;\nY;Aylık;;1;40;\nZ;Aylık;13;1;5;\n", encoding="utf-8")
            yy.write_text("Mükellef;Beyanname\nA;KDV\nA;kdv\nA;Özel\nA;X\n", encoding="utf-8")
            dy.write_text("Mükellef;Beyanname;Dönem;Verildi\nA;KDV;09.2026;20.10.2026\nA;KDV;eylül;\nA;KDV;2026-08;dün\n", encoding="utf-8")
            turler, yukler, durum, tatil, uy = main.oku(ty, yy, dy)
            self.assertEqual((turler["kdv"]["sonra"], turler["ozel"]["aylar"], turler["ozel"]["beyan"], sorted(turler), [(y["mukellef"], y["tur"]) for y in yukler], durum, tatil, sorted(u["tur"] for u in uy)),
                             (1, {6}, "son", ["kdv", "ozel"], [("A", "kdv"), ("A", "ozel")], {("a", "kdv", (2026, 9)): {"verildi": date(2026, 10, 20), "odendi": None}}, set(),
                              ["Beyanname türü okunamadı"] * 3 + ["Durum satırı okunamadı"] * 2 + ["Tanımsız beyanname türü"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--yukumlulukler", str(yy), "--turler", str(ty), "--cikti", c]
            self.assertEqual(main.main(arg + ["--durum", str(dy), "--ay", "2026-10", "--tarih", "26.10.2026", "--yaklasan", "5"]), 0)
            self.assertEqual(main.main(arg + ["--ay", "2026-08", "--tarih", "26.10.2026"]), 0)                  # Özel: Haziran + 2 ay = Ağustos
            for ek in (["--ay", "x"], ["--tarih", "x"], ["--yaklasan", "-1"]):
                self.assertEqual(main.main(arg + ek), 2)
            self.assertEqual(main.main(["--yukumlulukler", str(yy), "--cikti", c]), 2)
            self.assertEqual(main.main(arg + ["--durum", str(Path(tmp) / "yok.csv")]), 1)
            yy.write_text("Mükellef;Beyanname\nA;Özel\n", encoding="utf-8")
            self.assertEqual(main.main(arg + ["--ay", "2026-10"]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
