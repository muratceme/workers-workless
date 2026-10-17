"""
Yedek Parça Sipariş Önerisi — Gerisi Bende kod bloğu
Yedek Parça › Yedek Parça Danışmanı

Parça bazında tüketim, bekleyen iş emirleri ve tedarik süresine göre stok siparişi önerir.
  Günlük tüketim      = aylık ortalama tüketim ÷ --ay-gun (30)
  Kullanılabilir stok = eldeki stok + yoldaki sipariş − bekleyen iş emri ihtiyacı
  Sipariş noktası     = günlük tüketim × (tedarik süresi + --emniyet-gun 7)
  Hedef stok          = günlük tüketim × (tedarik süresi + emniyet günü + --kapsama-gun 30)
  Önerilen sipariş    = kullanılabilir stok ≤ sipariş noktası ise (hedef stok − kullanılabilir stok), tam sayıya ve paket adedinin katına yukarı yuvarlanır
  Stok gün            = eldeki stok ÷ günlük tüketim
  Durum               = Acil (bekleyen iş emri karşılanamıyor veya eldeki stok tedarik süresince yetmiyor) · Sipariş ver · Yeterli · Hareketsiz (tüketim yok, stok var)
Ortalama tüketimi siz verirsiniz; kod mevsimselliği, kampanyayı ve model değişimini bilmez. Emniyet ve kapsama günleri varsayımdır, parça grubuna göre değiştirin.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
from decimal import ROUND_CEILING

SUTUN = {"no": ("parca no", "parca", "parca numarasi", "stok kodu"), "ad": ("parca adi", "ad", "aciklama", "tanim"), "stok": ("stok", "eldeki stok", "mevcut stok"), "yol": ("yoldaki", "yoldaki siparis", "acik siparis"),
         "tuketim": ("aylik tuketim", "aylik ortalama tuketim", "ortalama tuketim", "tuketim"), "bekleyen": ("bekleyen is emri", "bekleyen ihtiyac", "rezerve", "bekleyen"), "tedarik": ("tedarik suresi", "tedarik suresi gun", "temin suresi"),
         "paket": ("paket adedi", "paket", "asgari siparis"), "fiyat": ("birim fiyat", "fiyat", "alis fiyati")}
Z = Decimal(0)
DURUM_RENK = {"Acil": KIRMIZI, "Sipariş ver": SARI, "Yeterli": YESIL, "Hareketsiz": MAVI}


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("no", "stok", "tuketim")):
        no = metin(r["no"])
        if not no:
            continue
        d = {a: para(r[a]) if metin(r[a]) else v for a, v in (("stok", None), ("yol", Z), ("tuketim", None), ("bekleyen", Z), ("tedarik", Z), ("paket", Decimal(1)), ("fiyat", None))}
        bozuk = [a for a in ("stok", "tuketim") if d[a] is None] + [a for a in ("yol", "bekleyen", "tedarik", "paket") if d[a] is None] + [a for a, v in d.items() if v is not None and v < 0] + (["paket"] if d["paket"] == 0 else [])
        if bozuk:
            uy.append(uyari("Yüksek", "Parça satırı okunamadı", f"{no} (satır {r['_satir']})", f"Okunamayan veya negatif alan: {', '.join(dict.fromkeys(bozuk))}. Parça değerlendirilmedi."))
            continue
        if katla(no) in gorulen:
            uy.append(uyari("Orta", "Mükerrer parça", no, f"Parça birden fazla satırda (satır {r['_satir']}). İlk satır kullanıldı; depolar ayrıysa stokları toplayıp tek satır girin."))
            continue
        gorulen.add(katla(no))
        l.append({"no": no, "ad": metin(r["ad"])} | d)
    return l, uy


def oner(parcalar: list[dict], emniyet_gun: int = 7, kapsama_gun: int = 30, ay_gun: int = 30) -> dict:
    if not parcalar:
        raise ValueError("Değerlendirilecek parça yok.")
    uy, liste = [], []
    for p in parcalar:
        gunluk = p["tuketim"] / ay_gun
        kullanilabilir = p["stok"] + p["yol"] - p["bekleyen"]
        nokta, hedef = gunluk * (p["tedarik"] + emniyet_gun), gunluk * (p["tedarik"] + emniyet_gun + kapsama_gun)
        ihtiyac = hedef - kullanilabilir if kullanilabilir <= nokta and (gunluk or kullanilabilir < 0) else Z
        adet = (max(ihtiyac, Z).quantize(Decimal(1), rounding=ROUND_CEILING) / p["paket"]).quantize(Decimal(1), rounding=ROUND_CEILING) * p["paket"]
        if p["bekleyen"] > p["stok"] + p["yol"]:
            durum, neden = "Acil", f"Bekleyen iş emri ({p['bekleyen']:f}) eldeki ve yoldaki stoktan ({p['stok'] + p['yol']:f}) fazla"
        elif gunluk and p["stok"] - p["bekleyen"] < gunluk * p["tedarik"]:
            durum, neden = "Acil", f"Eldeki stok tedarik süresindeki ({p['tedarik']:f} gün) tüketimi karşılamıyor"
        elif adet:
            durum, neden = "Sipariş ver", "Kullanılabilir stok sipariş noktasının altında"
        elif not gunluk and p["stok"]:
            durum, neden = "Hareketsiz", "Tüketim yok, stokta ürün var"
        else:
            durum, neden = "Yeterli", ""
        if durum == "Acil":
            uy.append(uyari("Yüksek", "Acil sipariş", p["no"], f"{p['ad']}: {neden}. Önerilen sipariş {adet:f} adet" + (f"; yolda {p['yol']:f} adet var, teslim tarihini teyit edin." if p["yol"] else ".")))
        elif durum == "Hareketsiz":
            uy.append(uyari("Bilgi", "Hareketsiz stok", p["no"], f"{p['ad']}: {p['stok']:f} adet stok var, tüketim yok" + (f" (değeri {p['stok'] * p['fiyat']:,.2f})." if p["fiyat"] else ".")))
        liste.append(p | {"gunluk": yuvarla(gunluk), "kullanilabilir": kullanilabilir, "nokta": yuvarla(nokta), "hedef": yuvarla(hedef), "stok_gun": yuvarla(p["stok"] / gunluk, "0.1") if gunluk else None, "oneri": adet,
                          "tutar": yuvarla(adet * p["fiyat"]) if p["fiyat"] is not None else None, "durum": durum, "neden": neden})
    sira = {"Acil": 0, "Sipariş ver": 1, "Yeterli": 2, "Hareketsiz": 3}
    liste.sort(key=lambda x: (sira[x["durum"]], x["no"]))
    say = {d: sum(1 for x in liste if x["durum"] == d) for d in sira}
    return {"parcalar": liste, "say": say, "toplam": sum((x["tutar"] or Z for x in liste), Z), "kalem": sum(1 for x in liste if x["oneri"]), "fiyatsiz": sum(1 for x in liste if x["oneri"] and x["tutar"] is None),
            "ayar": (emniyet_gun, kapsama_gun, ay_gun), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    eg, kg, ag = r["ayar"]
    ws = wb.active
    ws.title = "Sipariş Önerisi"
    _baslik(ws, ["Durum", "Parça No", "Parça Adı", "Eldeki Stok", "Yoldaki", "Bekleyen İş Emri", "Kullanılabilir", "Aylık Tüketim", "Stok Gün", "Tedarik (gün)", "Sipariş Noktası", "Hedef Stok", "Önerilen Sipariş", "Paket", "Birim Fiyat",
                 "Sipariş Tutarı", "Açıklama", "Onay (siz doldurun)"], (12, 16, 30, 10, 9, 10, 11, 10, 9, 9, 11, 10, 11, 7, 12, 14, 50, 20), "D2")
    for p in r["parcalar"]:
        _satir(ws, [p["durum"], p["no"], p["ad"], p["stok"], p["yol"] or None, p["bekleyen"] or None, p["kullanilabilir"], p["tuketim"], p["stok_gun"], p["tedarik"], p["nokta"], p["hedef"], p["oneri"] or None, p["paket"], p["fiyat"], p["tutar"] if p["oneri"] else None,
                    p["neden"], ""], para_s=(11, 12, 15, 16), dolgu={1: DURUM_RENK[p["durum"]], 13: SARI if p["oneri"] else None}, sar=(17,))
    n = _satir(ws, ["Toplam", "", f"{r['kalem']} kalem", None, None, None, None, None, None, None, None, None, None, None, None, r["toplam"]], para_s=(16,))
    for c in ws[n]:
        c.font = Font(bold=True)
    ws.append([])
    ws.append([f"Günlük tüketim = aylık ÷ {ag}. Sipariş noktası = günlük × (tedarik + {eg} emniyet günü); hedef stok buna {kg} günlük kapsama ekler. Öneri paket adedinin katına yukarı yuvarlanır."
               + (f" {r['fiyatsiz']} kalemin fiyatı olmadığı için toplam eksiktir." if r["fiyatsiz"] else "")])
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, **ayar) -> dict:
    parcalar, uy = oku(yol)
    r = oner(parcalar, **ayar)
    uy += r["uyarilar"]
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Tüketim, bekleyen iş emri ve tedarik süresine göre yedek parça siparişi önerir.")
    p.add_argument("--parcalar", type=Path, default=ORNEK / "parcalar.csv", help="Parça No, Parça Adı, Stok, Yoldaki, Aylık Tüketim, Bekleyen İş Emri, Tedarik Süresi (gün), Paket Adedi, Birim Fiyat")
    p.add_argument("--emniyet-gun", type=int, default=7, help="Emniyet stoku, gün (varsayılan 7)")
    p.add_argument("--kapsama-gun", type=int, default=30, help="Siparişin karşılayacağı ek gün (varsayılan 30)")
    p.add_argument("--ay-gun", type=int, default=30, help="Aylık tüketimi güne çevirmek için ay gün sayısı (varsayılan 30)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "yedek_parca_siparis_onerisi.xlsx")
    a = p.parse_args(argv)
    if a.emniyet_gun < 0 or a.kapsama_gun <= 0 or a.ay_gun <= 0:
        print("[X] --emniyet-gun sıfır veya pozitif; --kapsama-gun ve --ay-gun pozitif olmalı")
        return 2
    if not a.parcalar.exists():
        print(f"[X] Dosya bulunamadı: {a.parcalar}")
        return 1
    try:
        r = calistir(a.parcalar, a.cikti, emniyet_gun=a.emniyet_gun, kapsama_gun=a.kapsama_gun, ay_gun=a.ay_gun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    s = r["say"]
    print(f"[OK] {len(r['parcalar'])} parça · {s['Acil']} acil · {s['Sipariş ver']} sipariş ver · {s['Yeterli']} yeterli · {s['Hareketsiz']} hareketsiz · {r['kalem']} kalem sipariş, toplam {r['toplam']:,.2f}")
    for x in r["parcalar"]:
        if x["durum"] == "Acil":
            print(f"[!] {x['no']} {x['ad']}: {x['oneri']:f} adet · {x['neden']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
