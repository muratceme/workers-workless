import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def p(no, stok, tuketim, yol=0, bekleyen=0, tedarik=0, paket=1, fiyat=None):
    return {"no": no, "ad": "A", "stok": D(stok), "yol": D(yol), "tuketim": D(tuketim), "bekleyen": D(bekleyen), "tedarik": D(tedarik), "paket": D(paket), "fiyat": None if fiyat is None else D(fiyat)}


def tek(*a, **kw):
    ayar = {k: kw.pop(k) for k in list(kw) if k.endswith("_gun")}
    x = main.oner([p("P", *a, **kw)], **ayar)["parcalar"][0]
    return x["durum"], x["oneri"]


class OneriTesti(unittest.TestCase):
    def test_hesap(self):
        x = main.oner([p("P1", 20, 60, bekleyen=5, tedarik=10, paket=10, fiyat=100)])["parcalar"][0]
        self.assertEqual((x["gunluk"], x["kullanilabilir"], x["nokta"], x["hedef"], x["stok_gun"], x["oneri"], x["tutar"], x["durum"]), (D("2.00"), D(15), D("34.00"), D("94.00"), D("10.0"), D(80), D("8000.00"), "Acil"))
        x = main.oner([p("P6", 2, 10, tedarik=3, paket=5)])["parcalar"][0]
        self.assertEqual((x["gunluk"], x["nokta"], x["hedef"], x["stok_gun"], x["oneri"], x["tutar"], x["durum"]), (D("0.33"), D("3.33"), D("13.33"), D("6.0"), D(15), None, "Sipariş ver"))

    def test_durumlar(self):
        self.assertEqual((tek(100, 30, tedarik=5), tek(18, 30, yol=2, tedarik=15), tek(5, 30, yol=15, tedarik=15), tek(18, 30, yol=5, tedarik=15), tek(7, 0, tedarik=10), tek(1, 0, bekleyen=3, tedarik=20), tek(0, 0)),
                         (("Yeterli", D(0)), ("Sipariş ver", D(32)), ("Acil", D(32)), ("Yeterli", D(0)), ("Hareketsiz", D(0)), ("Acil", D(2)), ("Yeterli", D(0))))
        self.assertEqual((tek(12, 30, tedarik=5), tek(13, 30, tedarik=5), tek(12, 30, tedarik=5, emniyet_gun=0), tek(12, 30, tedarik=5, kapsama_gun=60), tek(12, 31, tedarik=5, ay_gun=31)),
                         (("Sipariş ver", D(30)), ("Yeterli", D(0)), ("Yeterli", D(0)), ("Sipariş ver", D(60)), ("Sipariş ver", D(30))))
        self.assertEqual((tek(4, 30, tedarik=5), tek(10, 30, bekleyen=6, tedarik=5), tek(10, 30, yol=10, bekleyen=25, tedarik=5)), (("Acil", D(38)), ("Acil", D(38)), ("Acil", D(47))))

    def test_siralama_ve_toplam(self):
        r = main.oner([p("Z", 100, 30, tedarik=5, fiyat=10), p("B", 12, 30, tedarik=5, fiyat="2.50"), p("A", 4, 30, tedarik=5), p("H", 7, 0, fiyat=20), p("C", 1, 0, bekleyen=3, fiyat=1000)])
        self.assertEqual(([(x["no"], x["durum"]) for x in r["parcalar"]], r["say"], r["toplam"], r["kalem"], r["fiyatsiz"], sorted((u["tur"], u["kim"]) for u in r["uyarilar"])),
                         ([("A", "Acil"), ("C", "Acil"), ("B", "Sipariş ver"), ("Z", "Yeterli"), ("H", "Hareketsiz")], {"Acil": 2, "Sipariş ver": 1, "Yeterli": 1, "Hareketsiz": 1}, D("2075.00"), 3, 1,
                          [("Acil sipariş", "A"), ("Acil sipariş", "C"), ("Hareketsiz stok", "H")]))
        with self.assertRaises(ValueError):
            main.oner([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "y.xlsx"
            r = main.calistir(ORNEK / "parcalar.csv", cikti)
            self.assertEqual([(x["no"], x["durum"], x["oneri"], x["tutar"]) for x in r["parcalar"]],
                             [("YP-1001", "Acil", D(80), D("8000.00")), ("YP-1005", "Acil", D(2), D("17000.00")), ("YP-1003", "Sipariş ver", D(32), D("30400.00")), ("YP-1006", "Sipariş ver", D(15), None),
                              ("YP-1002", "Yeterli", D(0), D("0.00")), ("YP-1004", "Hareketsiz", D(0), D("0.00"))])
            self.assertEqual((r["toplam"], r["kalem"], r["fiyatsiz"]), (D("55400.00"), 4, 1))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Sipariş Önerisi", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "p.csv"
            y.write_text("Stok Kodu;Tanım;Mevcut Stok;Tüketim;Paket\nA;x;1.250;30;\nA;x;5;5;\nB;y;;5;\nC;z;5;-1;\nD;w;5;5;0\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["stok"], x["paket"], x["yol"], x["fiyat"]) for x in l], sorted(u["tur"] for u in uy)), ([("A", D(1250), D(1), D(0), None)], ["Mükerrer parça"] + ["Parça satırı okunamadı"] * 3))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--parcalar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--emniyet-gun", "10", "--kapsama-gun", "45", "--ay-gun", "26"]), 0)
            for ek in (["--emniyet-gun", "-1"], ["--kapsama-gun", "0"], ["--ay-gun", "0"]):
                self.assertEqual(main.main(arg + ek), 2)
            self.assertEqual(main.main(["--parcalar", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Parça No;Stok\nA;1\n", "Parça No;Stok;Aylık Tüketim\nA;x;1\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
