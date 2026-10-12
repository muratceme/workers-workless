"""
Numune Ölçü Kontrolü — Workers / Workless kod bloğu
Modelhane › Modelist

Numune ölçümlerini ölçü tablosu ve toleranslarla karşılaştırır; tolerans dışı ölçü noktalarını raporlar.
  Fark          = ölçülen − ölçü tablosundaki beden değeri
  Tolerans içi  = −tolerans (−) ≤ fark ≤ +tolerans (+)        (sınır değer tolerans içindedir)
  Düzeltme      = −fark  (ör. fark +2,0 ise "2,0 küçült")
  Numune sonucu = kritik bir ölçü tolerans dışındaysa "Kritik ölçü tolerans dışı"; yalnız kritik olmayanlar dışındaysa "Tolerans dışı"; yoksa "Uygun"
Ölçü tablosu geniş biçimdedir: her beden bir sütun. Sayılarda virgül ve nokta ondalık ayırıcıdır; birim (cm / inç) tabloyla aynı olmalıdır.
Kod ölçüm yöntemini (yarım / tam ölçü, ölçü alma noktası) bilmez; onay kararı müşteri ve modelhaneye aittir.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SABIT = {"model": ("model", "model no", "model kodu", "stil"), "nokta": ("olcu noktasi", "olcu", "olcu yeri", "pom"), "kritik": ("kritik", "kritik olcu"),
         "tol": ("tolerans", "tolerans +", "tol"), "tol_eksi": ("tolerans eksi", "eksi tolerans", "tol eksi"), "tol_arti": ("tolerans arti", "arti tolerans", "tol arti")}
OLC_SUTUN = {"model": ("model", "model no", "model kodu", "stil"), "beden": ("beden",), "numune": ("numune", "numune no", "numune turu"), "nokta": ("olcu noktasi", "olcu", "olcu yeri", "pom"),
             "deger": ("olculen", "olculen deger", "olcum", "numune olcusu")}
EVET = ("evet", "e", "x", "1", "kritik", "var")


def sayi(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    try:
        return Decimal(re.sub(r"[^\d,.\-]", "", str(x).replace("−", "-")).replace(",", "."))
    except InvalidOperation:
        return None


def _baslik_katla(c) -> str:
    s = metin(c).replace("±", " ").replace("(-)", " eksi ").replace("(+)", " arti ").replace("(−)", " eksi ")
    return katla(s)


def tablo_yukle(yol: Path) -> tuple[dict, list[str], list[dict]]:
    """(tablo[(model, nokta)] = {beden: değer, ...}, beden adları, uyarılar)"""
    s, uy = tablo_oku(yol), []
    for bi, r in enumerate(s[:15]):
        b = [_baslik_katla(c) for c in r]
        es = {k: next((b.index(a) for a in adlar if a in b), None) for k, adlar in SABIT.items()}
        if es["nokta"] is not None:
            break
    else:
        raise ValueError(f"{yol.name}: başlık satırı bulunamadı; 'Ölçü Noktası' sütunu gerekli.")
    bedenler = [(j, metin(c)) for j, c in enumerate(s[bi]) if metin(c) and j not in es.values()]
    if not bedenler:
        raise ValueError(f"{yol.name}: beden sütunu yok.")
    tablo = {}
    for n, r in enumerate(s[bi + 1:], bi + 2):
        al = lambda k: r[es[k]] if es[k] is not None and es[k] < len(r) else None   # noqa: E731
        nokta = metin(al("nokta"))
        if not nokta:
            continue
        t = sayi(al("tol"))
        eksi, arti = sayi(al("tol_eksi")), sayi(al("tol_arti"))
        eksi, arti = abs(eksi) if eksi is not None else (abs(t) if t is not None else None), abs(arti) if arti is not None else (abs(t) if t is not None else None)
        if eksi is None or arti is None:
            uy.append(uyari("Yüksek", "Tolerans okunamadı", nokta, f"{yol.name} satır {n}: tolerans boş veya okunamadı. Ölçü noktası kontrol edilemedi."))
            continue
        d = {}
        for j, ad in bedenler:
            v = sayi(r[j] if j < len(r) else None)
            if v is None and j < len(r) and metin(r[j]):
                uy.append(uyari("Orta", "Tablo değeri okunamadı", f"{nokta} · {ad}", f"{yol.name} satır {n}: '{metin(r[j])}'."))
            if v is not None:
                d[katla(ad)] = v
        tablo[(katla(al("model")), katla(nokta))] = {"nokta": nokta, "kritik": katla(al("kritik")) in EVET, "eksi": eksi, "arti": arti, "beden": d}
    if not tablo:
        raise ValueError(f"{yol.name}: kullanılabilir ölçü noktası yok.")
    return tablo, [ad for _, ad in bedenler], uy


def olcum_oku(yol: Path) -> tuple[list[dict], list[dict]]:
    l, uy = [], []
    for r in kayitlar(yol, OLC_SUTUN, ("beden", "nokta", "deger")):
        if not metin(r["nokta"]):
            continue
        v = sayi(r["deger"])
        if v is None:
            uy.append(uyari("Yüksek", "Ölçüm okunamadı", f"{metin(r['beden'])} · {metin(r['nokta'])}", f"{yol.name} satır {r['_satir']}: '{metin(r['deger'])}'. Ölçüm alınmadı."))
            continue
        l.append({"model": metin(r["model"]), "beden": metin(r["beden"]), "numune": metin(r["numune"]), "nokta": metin(r["nokta"]), "deger": v})
    return l, uy


def kontrol(olcumler: list[dict], tablo: dict) -> dict:
    uy, sonuc, yok = [], [], defaultdict(int)
    for o in olcumler:
        t = tablo.get((katla(o["model"]), katla(o["nokta"]))) or tablo.get(("", katla(o["nokta"])))
        hedef = t["beden"].get(katla(o["beden"])) if t else None
        if hedef is None:
            yok[(o["nokta"], o["beden"], bool(t))] += 1
            continue
        fark = o["deger"] - hedef
        durum = "Tolerans içinde" if -t["eksi"] <= fark <= t["arti"] else "Büyük" if fark > 0 else "Küçük"
        sonuc.append(o | {"hedef": hedef, "fark": fark, "eksi": t["eksi"], "arti": t["arti"], "kritik": t["kritik"], "durum": durum,
                          "duzeltme": "" if durum == "Tolerans içinde" else f"{abs(fark).normalize():f} {'küçült' if fark > 0 else 'büyüt'}".replace(".", ",")})
    for (nokta, beden, var), n in yok.items():
        uy.append(uyari("Yüksek", "Ölçü tabloda yok", f"{nokta} · {beden}", f"{n} ölçüm için " + ("tabloda bu beden değeri yok." if var else "ölçü tablosunda bu ölçü noktası yok.") + " Kontrol edilmedi."))
    numune = {}
    for s in sonuc:
        x = numune.setdefault((s["model"], s["beden"], s["numune"]), {"model": s["model"], "beden": s["beden"], "numune": s["numune"], "olcu": 0, "kritik_dis": [], "dis": []})
        x["olcu"] += 1
        if s["durum"] != "Tolerans içinde":
            x["kritik_dis" if s["kritik"] else "dis"].append(s["nokta"])
    for x in numune.values():
        x["sonuc"] = "Kritik ölçü tolerans dışı" if x["kritik_dis"] else "Tolerans dışı" if x["dis"] else "Uygun"
    nokta = {}
    for s in sonuc:
        x = nokta.setdefault(katla(s["nokta"]), {"nokta": s["nokta"], "kritik": s["kritik"], "olcu": 0, "dis": 0, "buyuk": 0, "kucuk": 0, "farklar": []})
        x["olcu"] += 1
        x["farklar"].append(s["fark"])
        x["dis"] += s["durum"] != "Tolerans içinde"
        x["buyuk"] += s["durum"] == "Büyük"
        x["kucuk"] += s["durum"] == "Küçük"
    for x in nokta.values():
        x["ort_fark"] = yuvarla(sum(x["farklar"]) / len(x["farklar"]))
    return {"sonuc": sonuc, "numune": list(numune.values()), "nokta": sorted(nokta.values(), key=lambda x: (-x["dis"], not x["kritik"], x["nokta"])), "uyarilar": uy,
            "dis": sum(s["durum"] != "Tolerans içinde" for s in sonuc)}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    renk = {"Uygun": YESIL, "Tolerans dışı": SARI, "Kritik ölçü tolerans dışı": KIRMIZI}
    ws = wb.active
    ws.title = "Numune Özeti"
    _baslik(ws, ["Model", "Beden", "Numune", "Ölçü Sayısı", "Sonuç", "Tolerans Dışı Kritik Ölçüler", "Tolerans Dışı Diğer Ölçüler", "Karar / Yorum (siz doldurun)"], (16, 9, 14, 9, 26, 40, 40, 34))
    for x in sorted(r["numune"], key=lambda x: list(renk).index(x["sonuc"]), reverse=True):
        _satir(ws, [x["model"], x["beden"], x["numune"], x["olcu"], x["sonuc"], ", ".join(x["kritik_dis"]), ", ".join(x["dis"]), ""], dolgu={5: renk[x["sonuc"]]}, sar=(6, 7))
    ws = wb.create_sheet("Ölçü Karşılaştırma")
    _baslik(ws, ["Model", "Beden", "Numune", "Ölçü Noktası", "Kritik", "Tablo Değeri", "Ölçülen", "Fark", "Tolerans (−)", "Tolerans (+)", "Durum", "Düzeltme"], (16, 9, 14, 30, 8, 11, 11, 9, 11, 11, 16, 16))
    for s in sorted(r["sonuc"], key=lambda s: (s["durum"] == "Tolerans içinde", not s["kritik"])):
        _satir(ws, [s["model"], s["beden"], s["numune"], s["nokta"], "Evet" if s["kritik"] else "", s["hedef"], s["deger"], s["fark"], s["eksi"], s["arti"], s["durum"], s["duzeltme"]],
               dolgu={11: None if s["durum"] == "Tolerans içinde" else KIRMIZI if s["kritik"] else SARI})
    ws = wb.create_sheet("Ölçü Noktası Özeti")
    _baslik(ws, ["Ölçü Noktası", "Kritik", "Ölçüm Sayısı", "Tolerans Dışı", "Büyük", "Küçük", "Ortalama Fark"], (30, 8, 10, 10, 8, 8, 12))
    for x in r["nokta"]:
        _satir(ws, [x["nokta"], "Evet" if x["kritik"] else "", x["olcu"], x["dis"] or None, x["buyuk"] or None, x["kucuk"] or None, x["ort_fark"]], para_s=(7,), dolgu={4: SARI if x["dis"] else None})
    ws.append([])
    ws.append(["Aynı ölçü noktası tüm bedenlerde aynı yönde sapıyorsa neden kalıp veya çekme payı olabilir; tek bedende sapıyorsa serileme (grading) kontrol edilmelidir."])
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(olcum_yol: Path, tablo_yol: Path, cikti: Path) -> dict:
    tablo, bedenler, uy1 = tablo_yukle(tablo_yol)
    olcumler, uy2 = olcum_oku(olcum_yol)
    if not olcumler:
        raise ValueError(f"{olcum_yol.name}: okunabilen ölçüm yok.")
    r = kontrol(olcumler, tablo)
    r["uyarilar"], r["bedenler"] = sirala(uy1 + uy2 + r["uyarilar"]), bedenler
    rapor_yaz(cikti, r, r["uyarilar"])
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Numune ölçümlerini ölçü tablosu ve toleranslarla karşılaştırır.")
    p.add_argument("--olcumler", type=Path, default=ORNEK / "numune_olcumleri.csv", help="Model, Beden, Numune, Ölçü Noktası, Ölçülen")
    p.add_argument("--tablo", type=Path, default=ORNEK / "olcu_tablosu.csv", help="Ölçü Noktası, Kritik, Tolerans (±) veya Tolerans (-) ve Tolerans (+), sonra her beden bir sütun")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "numune_olcu_kontrolu.xlsx")
    a = p.parse_args(argv)
    for y in (a.olcumler, a.tablo):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.olcumler, a.tablo, a.cikti)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(r['numune'])} numune · {len(r['sonuc'])} ölçü · {r['dis']} tolerans dışı")
    for x in r["numune"]:
        if x["sonuc"] != "Uygun":
            print(f"[!] {x['model']} {x['beden']} {x['numune']}: {x['sonuc']} · " + ", ".join(x["kritik_dis"] + x["dis"]))
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
