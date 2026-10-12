import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def tb(nokta, kritik, eksi, arti, **beden):
    return {"nokta": nokta, "kritik": kritik, "eksi": D(str(eksi)), "arti": D(str(arti)), "beden": {k.lower(): D(str(v)) for k, v in beden.items()}}


TABLO = {("", "gogus"): tb("Göğüs", True, 1, 1, S=48, M=51), ("", "kol"): tb("Kol", False, 1, 0.5, S=20, M=21), ("x1", "boy"): tb("Boy", True, 1.5, 1.5, M=70)}


def ol(nokta, deger, beden="M", model="", numune="N"):
    return {"model": model, "beden": beden, "numune": numune, "nokta": nokta, "deger": D(str(deger))}


class KontrolTesti(unittest.TestCase):
    def test_tolerans_sinirlari(self):
        r = main.kontrol([ol("Göğüs", v) for v in (50, 49.9, 52, 52.1)] + [ol("Kol", v) for v in (20, 19.9, 21.5, 21.6)], TABLO)
        self.assertEqual([(s["fark"], s["durum"], s["duzeltme"]) for s in r["sonuc"]],
                         [(D(-1), "Tolerans içinde", ""), (D("-1.1"), "Küçük", "1,1 büyüt"), (D(1), "Tolerans içinde", ""), (D("1.1"), "Büyük", "1,1 küçült"),
                          (D(-1), "Tolerans içinde", ""), (D("-1.1"), "Küçük", "1,1 büyüt"), (D("0.5"), "Tolerans içinde", ""), (D("0.6"), "Büyük", "0,6 küçült")])   # eksi ve artı tolerans farklı
        self.assertEqual(r["dis"], 4)

    def test_numune_sonucu(self):
        r = main.kontrol([ol("Göğüs", 51, "M", numune="A"), ol("Kol", 21, "M", numune="A"), ol("Göğüs", 48, "S", numune="B"), ol("Kol", 23, "S", numune="B"),
                          ol("Göğüs", 53, "M", numune="C"), ol("Kol", 25, "M", numune="C")], TABLO)
        self.assertEqual({x["numune"]: (x["sonuc"], x["kritik_dis"], x["dis"]) for x in r["numune"]},
                         {"A": ("Uygun", [], []), "B": ("Tolerans dışı", [], ["Kol"]), "C": ("Kritik ölçü tolerans dışı", ["Göğüs"], ["Kol"])})
        n = {x["nokta"]: (x["olcu"], x["dis"], x["buyuk"], x["kucuk"], x["ort_fark"]) for x in r["nokta"]}
        self.assertEqual(n, {"Kol": (3, 2, 2, 0, D("2.33")), "Göğüs": (3, 1, 1, 0, D("0.67"))})

    def test_model_ve_eksik_tablo(self):
        r = main.kontrol([ol("Boy", 70, model="X1"), ol("Boy", 70, model="Y2"), ol("Göğüs", 51, "XL"), ol("Göğüs", 51, model="Y2")], TABLO)
        self.assertEqual([(s["model"], s["nokta"]) for s in r["sonuc"]], [("X1", "Boy"), ("Y2", "Göğüs")])            # modelsiz tablo satırı tüm modellere uygulanır
        self.assertEqual(sorted(u["aciklama"][:40] for u in r["uyarilar"]), ["1 ölçüm için tabloda bu beden değeri yok", "1 ölçüm için ölçü tablosunda bu ölçü nok"])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "n.xlsx"
            r = main.calistir(ORNEK / "numune_olcumleri.csv", ORNEK / "olcu_tablosu.csv", cikti)
            self.assertEqual({x["beden"]: (x["sonuc"], x["kritik_dis"], x["dis"]) for x in r["numune"]},
                             {"S": ("Uygun", [], []), "M": ("Kritik ölçü tolerans dışı", ["Boy (omuzdan)"], ["Kol boyu"]), "L": ("Tolerans dışı", [], ["Etek genişliği (1/2)"])})
            self.assertEqual((len(r["sonuc"]), r["dis"], r["bedenler"], r["uyarilar"]), (18, 3, ["S", "M", "L"], []))
            self.assertEqual(next(s["duzeltme"] for s in r["sonuc"] if s["beden"] == "M" and s["nokta"] == "Boy (omuzdan)"), "2 küçült")
            self.assertEqual(load_workbook(cikti).sheetnames, ["Numune Özeti", "Ölçü Karşılaştırma", "Ölçü Noktası Özeti", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            t, o = Path(tmp) / "t.csv", Path(tmp) / "o.csv"
            t.write_text("Model TS\nÖlçü Noktası;Kritik;Tolerans (-);Tolerans (+);S;M\nGöğüs;x;-1;0,5;48;51.5\nKol;;;;20;21\nBoy;;1;1;uzun;70\n", encoding="utf-8")
            tablo, bedenler, uy = main.tablo_yukle(t)
            g = tablo[("", "gogus")]
            self.assertEqual((sorted(k[1] for k in tablo), g["kritik"], g["eksi"], g["arti"], g["beden"], bedenler), (["boy", "gogus"], True, D(1), D("0.5"), {"s": D(48), "m": D("51.5")}, ["S", "M"]))
            self.assertEqual(sorted(u["tur"] for u in uy), ["Tablo değeri okunamadı", "Tolerans okunamadı"])
            o.write_text("Beden;Ölçü Noktası;Ölçülen\nM;Göğüs;52,0\nM;Boy;yok\n", encoding="utf-8")
            ol_, uy = main.olcum_oku(o)
            self.assertEqual(([x["deger"] for x in ol_], [u["tur"] for u in uy]), ([D("52.0")], ["Ölçüm okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            self.assertEqual(main.main(["--olcumler", str(o), "--tablo", str(t), "--cikti", c]), 0)
            self.assertEqual(main.main(["--olcumler", str(Path(tmp) / "yok.csv")]), 1)
            t.write_text("Ölçü Noktası;Tolerans\nGöğüs;1\n", encoding="utf-8")
            self.assertEqual(main.main(["--olcumler", str(o), "--tablo", str(t), "--cikti", c]), 1)       # beden sütunu yok
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
