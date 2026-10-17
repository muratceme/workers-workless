"""
Kontrol Testi Örneklem Seçimi — Gerisi Bende kod bloğu
İç Kontrol ve Risk Yönetimi › İç Kontrol Uzmanı

İşlem popülasyonundan risk ve tutar ağırlıklı, tekrarlanabilir (tohumlu) örneklem seçer.
  Tohum            = aynı dosya, aynı sıra, aynı tohum ve aynı parametrelerle her çalıştırmada aynı örneklem çıkar
  rastgele         = her kaleme tohumlu rastgele bir sayı verilir; en küçük sayıya sahip N kalem seçilir (eşit olasılık)
  tutar            = parasal birim örneklemesi: aralık = toplam tutar ÷ örneklem sayısı; tutarı aralığa eşit ya da büyük kalemler kesin seçilir,
                     kalanlar için aralık yeniden hesaplanır ve rastgele bir başlangıçtan itibaren her aralıkta bir kalem seçilir
                     (kalemin seçilme olasılığı tutarıyla orantılıdır)
  risk             = tabakalı: her risk tabakasına (Yüksek / Orta / Düşük) kalem sayısı × ağırlık oranında pay verilir; tabaka içinde rastgele seçilir
  Kapsama %        = örneklem ÷ popülasyon × 100   (adet ve tutar bazında)
Örneklem büyüklüğünü siz belirlersiniz. Kod örneklemi seçer; test sonucundan popülasyona hata oranı çıkarımı yapmaz.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
import random

SUTUN = {"no": ("islem no", "islem numarasi", "kayit no", "fis no", "no"), "tarih": ("tarih", "islem tarihi"), "tutar": ("tutar", "islem tutari"), "risk": ("risk", "risk seviyesi", "risk duzeyi"), "aciklama": ("aciklama", "islem turu", "tur")}
RISK = {"yuksek": "Yüksek", "orta": "Orta", "dusuk": "Düşük"}
TABAKALAR = ("Yüksek", "Orta", "Düşük", "Belirsiz")
Z, YUZ = Decimal(0), Decimal(100)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("no",)):
        no = metin(r["no"])
        if not no:
            continue
        tutar = para(r["tutar"])
        if metin(r["tutar"]) and tutar is None:
            uy.append(uyari("Yüksek", "Satır okunamadı", no, f"Satır {r['_satir']}: tutar '{metin(r['tutar'])}' okunamadı. Satır popülasyona alınmadı."))
            continue
        if katla(no) in gorulen:
            uy.append(uyari("Yüksek", "Mükerrer işlem numarası", no, f"Satır {r['_satir']}: aynı işlem numarası daha önce geçti; bu satır popülasyona alınmadı. Popülasyonun tam ve tekil olduğunu doğrulayın."))
            continue
        gorulen.add(katla(no))
        l.append({"no": no, "tarih": tarih(r["tarih"]), "tutar": tutar, "risk": RISK.get(katla(r["risk"]), "Belirsiz"), "aciklama": metin(r["aciklama"])})
    return l, uy


def dagit(adet: int, tabakalar: list[tuple[str, int, Decimal]]) -> dict[str, int]:
    """Örneklem sayısını tabakalara kalem sayısı × ağırlık oranında dağıtır (en büyük kalan yöntemi; tabaka kapasitesini aşmaz)."""
    sonuc, aktif, kalan = {ad: 0 for ad, _, _ in tabakalar}, [t for t in tabakalar if t[1] > 0], min(adet, sum(t[1] for t in tabakalar))
    while aktif and kalan > 0:
        skor = {ad: Decimal(n) * w for ad, n, w in aktif}
        if not sum(skor.values()):
            skor = {ad: Decimal(n) for ad, n, _ in aktif}
        toplam = sum(skor.values())
        ideal = {ad: Decimal(kalan) * skor[ad] / toplam for ad, _, _ in aktif}
        dolan = [t for t in aktif if ideal[t[0]] > t[1]]
        if dolan:
            for ad, n, _ in dolan:
                sonuc[ad] = n
                kalan -= n
            aktif = [t for t in aktif if t not in dolan]
            continue
        taban = {ad: int(ideal[ad]) for ad in ideal}
        artan = kalan - sum(taban.values())
        for ad in sorted(ideal, key=lambda a: (-(ideal[a] - taban[a]), [t[0] for t in aktif].index(a)))[:artan]:
            taban[ad] += 1
        sonuc.update(taban)
        break
    return sonuc


def sistematik(kalemler: list[dict], aralik: Decimal, baslangic: Decimal) -> list[dict]:
    """Kalemleri tutarlarıyla uç uca dizer; baslangic, baslangic + aralik, … noktalarına denk gelen kalemleri döndürür."""
    secilen, nokta, kum = [], baslangic, Z
    for k in kalemler:
        kum += k["tutar"]
        if nokta < kum:
            secilen.append(k)
            while nokta < kum:
                nokta += aralik
    return secilen


def sec(populasyon: list[dict], yontem: str = "rastgele", adet: int = 25, tohum: int = 2026, agirlik: tuple[Decimal, Decimal, Decimal] = (Decimal(3), Decimal(2), Decimal(1))) -> dict:
    if not populasyon:
        raise ValueError("Popülasyon boş.")
    uy, rng = [], random.Random(tohum)
    for k in populasyon:
        k["anahtar"], k["neden"] = rng.random(), None
    bilgi = {"yontem": yontem, "tohum": tohum, "istenen": adet}
    if adet >= len(populasyon):
        uy.append(uyari("Bilgi", "Örneklem popülasyonun tamamı", f"{len(populasyon)} kalem", f"İstenen örneklem sayısı ({adet}) popülasyondaki kalem sayısına eşit ya da büyük; kalemlerin tamamı seçildi."))
    if yontem == "rastgele":
        for k in sorted(populasyon, key=lambda k: k["anahtar"])[:adet]:
            k["neden"] = "Rastgele"
    elif yontem == "tutar":
        uygun = [k for k in populasyon if k["tutar"] is not None and k["tutar"] > 0]
        if len(uygun) < len(populasyon):
            uy.append(uyari("Orta", "Tutarı olmayan kalemler seçilemez", f"{len(populasyon) - len(uygun)} kalem", "Tutarı boş, sıfır ya da eksi olan kalemler tutar ağırlıklı yöntemde seçilemez; bunları ayrıca değerlendirin."))
        if not uygun:
            raise ValueError("Tutar ağırlıklı yöntem için tutarı pozitif kalem yok.")
        kalan, n = list(uygun), min(adet, len(uygun))
        while True:
            aralik = sum((k["tutar"] for k in kalan), Z) / (n - (len(uygun) - len(kalan))) if n > len(uygun) - len(kalan) else None
            buyuk = [k for k in kalan if aralik is not None and k["tutar"] >= aralik]
            if not buyuk:
                break
            for k in buyuk:
                k["neden"] = "Kesin seçim (tutar ≥ aralık)"
            kalan = [k for k in kalan if k["neden"] is None]
        if aralik is not None and kalan:
            baslangic = Decimal(str(rng.random())) * aralik
            for k in sistematik(kalan, aralik, baslangic):
                k["neden"] = "Tutar ağırlıklı (aralık)"
            bilgi |= {"aralik": yuvarla(aralik), "baslangic": yuvarla(baslangic)}
        bilgi["kesin"] = sum(1 for k in uygun if k["neden"] and k["neden"].startswith("Kesin"))
    else:
        if any(k["risk"] == "Belirsiz" for k in populasyon):
            uy.append(uyari("Bilgi", "Risk seviyesi belirsiz kalemler", f"{sum(1 for k in populasyon if k['risk'] == 'Belirsiz')} kalem", "Risk sütunu boş ya da tanınmayan kalemler 'Belirsiz' tabakasında, Orta ağırlığıyla değerlendirildi."))
        w = dict(zip(TABAKALAR, agirlik + (agirlik[1],)))
        pay = dagit(adet, [(t, sum(1 for k in populasyon if k["risk"] == t), w[t]) for t in TABAKALAR])
        for t in TABAKALAR:
            for k in sorted((k for k in populasyon if k["risk"] == t), key=lambda k: k["anahtar"])[:pay[t]]:
                k["neden"] = f"Risk tabakası: {t}"
        bilgi["agirlik"] = "; ".join(f"{t} {w[t]:f}" for t in TABAKALAR[:3])
    orneklem = [k for k in populasyon if k["neden"]]
    tutar = lambda l: sum((k["tutar"] or Z for k in l), Z)   # noqa: E731
    pt, ot = tutar(populasyon), tutar(orneklem)
    tabakalar = [{"tabaka": t, "populasyon": sum(1 for k in populasyon if k["risk"] == t), "orneklem": sum(1 for k in orneklem if k["risk"] == t), "populasyon_tutar": tutar([k for k in populasyon if k["risk"] == t]),
                  "orneklem_tutar": tutar([k for k in orneklem if k["risk"] == t])} for t in TABAKALAR if any(k["risk"] == t for k in populasyon)]
    return {"orneklem": orneklem, "populasyon": populasyon, "tabakalar": tabakalar, "bilgi": bilgi,
            "ozet": {"populasyon": len(populasyon), "orneklem": len(orneklem), "populasyon_tutar": pt, "orneklem_tutar": ot, "kapsama": yuvarla(Decimal(len(orneklem) * 100) / len(populasyon)), "kapsama_tutar": yuvarla(ot * YUZ / pt) if pt else None},
            "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, kaynak: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Örneklem"
    _baslik(ws, ["Sıra", "İşlem No", "Tarih", "Tutar", "Risk", "Açıklama", "Seçim Nedeni", "Test Sonucu", "Bulgu / Not", "Testi Yapan"], (6, 16, 12, 16, 10, 30, 30, 14, 40, 18))
    for i, k in enumerate(r["orneklem"], 1):
        _satir(ws, [i, k["no"], k["tarih"], k["tutar"], k["risk"], k["aciklama"], k["neden"], None, None, None], para_s=(4,), tarih_s=(3,), dolgu={7: SARI if k["neden"].startswith("Kesin") else None})
    ws = wb.create_sheet("Özet")
    o, b = r["ozet"], r["bilgi"]
    _baslik(ws, ["Gösterge", "Değer"], (40, 24))
    for etiket, deger, para_mi in (("Popülasyon (kalem)", o["populasyon"], False), ("Örneklem (kalem)", o["orneklem"], False), ("Kapsama % (adet)", o["kapsama"], True), ("Popülasyon tutarı", o["populasyon_tutar"], True),
                                   ("Örneklem tutarı", o["orneklem_tutar"], True), ("Kapsama % (tutar)", o["kapsama_tutar"], True)):
        _satir(ws, [etiket, deger], para_s=(2,) if para_mi else ())
    ws.append([])
    _baslik(ws, ["Seçim parametreleri (denetim izi)", "Değer"], dondur="A2")
    for etiket, deger in (("Kaynak dosya", kaynak), ("Yöntem", {"rastgele": "Rastgele (eşit olasılık)", "tutar": "Tutar ağırlıklı (parasal birim)", "risk": "Risk tabakalı"}[b["yontem"]]), ("Tohum", b["tohum"]), ("İstenen örneklem", b["istenen"]),
                          ("Aralık", b.get("aralik")), ("Rastgele başlangıç", b.get("baslangic")), ("Kesin seçilen kalem", b.get("kesin")), ("Tabaka ağırlıkları", b.get("agirlik"))):
        if deger is not None:
            _satir(ws, [etiket, float(deger) if isinstance(deger, Decimal) else deger], para_s=(2,) if isinstance(deger, Decimal) else ())
    ws.append([])
    _baslik(ws, ["Risk Tabakası", "Popülasyon", "Örneklem", "Popülasyon Tutarı", "Örneklem Tutarı"], dondur="A2")
    for t in r["tabakalar"]:
        _satir(ws, [t["tabaka"], t["populasyon"], t["orneklem"], t["populasyon_tutar"], t["orneklem_tutar"]], para_s=(4, 5))
    for s, g in zip("CDE", (12, 20, 20)):
        ws.column_dimensions[s].width = g
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, yontem: str = "rastgele", adet: int = 25, tohum: int = 2026, agirlik: tuple = (Decimal(3), Decimal(2), Decimal(1))) -> dict:
    populasyon, uy = oku(yol)
    r = sec(populasyon, yontem, adet, tohum, agirlik)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, yol.name)
    return r


def agirlik_coz(s: str) -> tuple[Decimal, ...] | None:
    d = tuple(para(x) for x in re.split(r"[;\s]+", s.strip()) if x)
    return d if len(d) == 3 and all(x is not None and x >= 0 for x in d) and any(d) else None


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="İşlem popülasyonundan tohumlu, tekrarlanabilir kontrol testi örneklemi seçer.")
    p.add_argument("--populasyon", type=Path, help="İşlem No, Tarih, Tutar, Risk (Yüksek / Orta / Düşük), Açıklama")
    p.add_argument("--yontem", choices=("rastgele", "tutar", "risk"), default="rastgele", help="rastgele: eşit olasılık (varsayılan); tutar: tutarla orantılı; risk: risk tabakalı")
    p.add_argument("--adet", type=int, help="Örneklem sayısı (varsayılan 25)")
    p.add_argument("--tohum", type=int, default=2026, help="Rastgelelik tohumu; çalışma kâğıdına yazın (varsayılan 2026)")
    p.add_argument("--agirlik", default="3;2;1", help="Risk yönteminde Yüksek;Orta;Düşük tabaka ağırlıkları (varsayılan 3;2;1)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kontrol_testi_orneklem.xlsx")
    a = p.parse_args(argv)
    ornek = a.populasyon is None
    yol = ORNEK / "populasyon.csv" if ornek else a.populasyon
    adet, agirlik = a.adet if a.adet is not None else 10 if ornek else 25, agirlik_coz(a.agirlik)
    if adet < 1 or agirlik is None:
        print("[X] --adet en az 1; --agirlik üç sayı olmalı (örn. 3;2;1), en az biri sıfırdan büyük")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, a.yontem, adet, a.tohum, agirlik)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] Popülasyon {o['populasyon']} kalem ({o['populasyon_tutar']:,.2f}) · örneklem {o['orneklem']} kalem ({o['orneklem_tutar']:,.2f}) · kapsama adet %{o['kapsama']}" + (f", tutar %{o['kapsama_tutar']}" if o["kapsama_tutar"] is not None else "")
          + f" · yöntem {a.yontem}, tohum {a.tohum}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
