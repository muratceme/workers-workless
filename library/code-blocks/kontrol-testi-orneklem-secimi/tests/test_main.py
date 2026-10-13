import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def pop(tutarlar, riskler=None):
    return [{"no": f"N{i}", "tarih": None, "tutar": None if t is None else D(str(t)), "risk": (riskler[i - 1] if riskler else "Orta"), "aciklama": ""} for i, t in enumerate(tutarlar, 1)]


def nolar(r):
    return [k["no"] for k in r["orneklem"]]


class SecimTesti(unittest.TestCase):
    def test_rastgele(self):
        a, b, c = (main.sec(pop(range(1, 101)), "rastgele", 15, t) for t in (7, 7, 8))
        self.assertEqual((nolar(a), len(nolar(a)), len(set(nolar(a))), {k["neden"] for k in a["orneklem"]}), (nolar(b), 15, 15, {"Rastgele"}))
        self.assertNotEqual(nolar(a), nolar(c))
        self.assertEqual(nolar(a), sorted(nolar(a), key=lambda n: int(n[1:])))
        self.assertEqual(a["ozet"], {"populasyon": 100, "orneklem": 15, "populasyon_tutar": D(5050), "orneklem_tutar": sum(k["tutar"] for k in a["orneklem"]), "kapsama": D("15.00"),
                                     "kapsama_tutar": (sum(k["tutar"] for k in a["orneklem"]) * 100 / D(5050)).quantize(D("0.01"))})
        r = main.sec(pop([5, None, 0]), "rastgele", 10, 1)
        self.assertEqual((nolar(r), r["ozet"]["kapsama"], r["ozet"]["populasyon_tutar"], [u["tur"] for u in r["uyarilar"]]), (["N1", "N2", "N3"], D("100.00"), D(5), ["Örneklem popülasyonun tamamı"]))
        self.assertEqual(main.sec(pop([None, None]), "rastgele", 1, 1)["ozet"]["kapsama_tutar"], None)
        with self.assertRaises(ValueError):
            main.sec([], "rastgele", 5, 1)

    def test_tutar(self):
        kalemler = pop([100, 50])
        self.assertEqual(([k["no"] for k in main.sistematik(kalemler, D(150), D(120))], [k["no"] for k in main.sistematik(kalemler, D(150), D(40))], [k["no"] for k in main.sistematik(kalemler, D(150), D(100))],
                          [k["no"] for k in main.sistematik(pop([10, 10, 10, 10, 10, 10]), D(20), D("9.99"))], [k["no"] for k in main.sistematik(pop([10, 10, 10, 10, 10, 10]), D(20), D(10))]),
                         (["N2"], ["N1"], ["N2"], ["N1", "N3", "N5"], ["N2", "N4", "N6"]))
        for tohum in (1, 2, 3):
            r = main.sec(pop([100, 300, 50, 250, 300]), "tutar", 4, tohum)
            self.assertEqual(([(k["no"], k["neden"]) for k in r["orneklem"] if k["neden"].startswith("Kesin")], len(r["orneklem"]), r["bilgi"]["kesin"], r["bilgi"]["aralik"]),
                             ([("N2", "Kesin seçim (tutar ≥ aralık)"), ("N4", "Kesin seçim (tutar ≥ aralık)"), ("N5", "Kesin seçim (tutar ≥ aralık)")], 4, 3, D("150.00")))
            self.assertTrue(D(0) <= r["bilgi"]["baslangic"] < D(150) and (("N1" in nolar(r)) == (r["bilgi"]["baslangic"] < 100)))
        buyuk = [12500, 8400, 150000, 3200, 47000, 9800, 22000, 5600, 310000, 14500, 7300, 68000, 2900, 18500, 95000, 4100, 26000, 11200, 530000, 6700, 39000, 15800, 8900, 72000, 3500, 21000, 125000, 9400, 44000, 17300]
        r1, r2 = main.sec(pop(buyuk), "tutar", 10, 2026), main.sec(pop(buyuk), "tutar", 10, 2026)
        self.assertEqual((nolar(r1), len(r1["orneklem"]), sorted(k["no"] for k in r1["orneklem"] if k["neden"].startswith("Kesin")), r1["bilgi"]["aralik"], r1["bilgi"]["kesin"]), (nolar(r2), 10, ["N19", "N27", "N3", "N9"], D("98933.33"), 4))
        r = main.sec(pop([10, -5, 0, None, 30, 60]), "tutar", 2, 5)
        self.assertEqual((len(r["orneklem"]), "N6" in nolar(r), {"N2", "N3", "N4"} & set(nolar(r)), [u["tur"] for u in r["uyarilar"]]), (2, True, set(), ["Tutarı olmayan kalemler seçilemez"]))
        self.assertEqual((nolar(main.sec(pop([1, 2, 3]), "tutar", 3, 1)), nolar(main.sec(pop([7, 7, 7]), "tutar", 9, 1))), (["N1", "N2", "N3"], ["N1", "N2", "N3"]))
        with self.assertRaises(ValueError):
            main.sec(pop([0, None]), "tutar", 1, 1)

    def test_risk(self):
        d = lambda adet, n, w: list(main.dagit(adet, [(t, x, D(y)) for t, x, y in zip(("Y", "O", "D"), n, w)]).values())   # noqa: E731
        self.assertEqual([d(15, (10, 30, 60), (3, 2, 1)), d(10, (10, 30, 60), (3, 2, 1)), d(12, (2, 30, 60), (10, 1, 1)), d(20, (1, 30, 60), (10, 1, 1)), d(10, (5, 10, 15), (3, 2, 1)), d(100, (5, 10, 15), (3, 2, 1)), d(6, (5, 0, 5), (1, 1, 1)),
                          d(3, (4, 4, 4), (0, 0, 0)), d(5, (10, 10, 10), (1, 0, 0)), d(7, (3, 10, 10), (1, 0, 0))],
                         [[3, 6, 6], [2, 4, 4], [2, 3, 7], [1, 6, 13], [3, 4, 3], [5, 10, 15], [3, 0, 3], [1, 1, 1], [5, 0, 0], [3, 2, 2]])
        riskler = ["Yüksek"] * 5 + ["Orta"] * 10 + ["Düşük"] * 15
        r, r2 = main.sec(pop(range(1, 31), riskler), "risk", 10, 3), main.sec(pop(range(1, 31), riskler), "risk", 10, 3)
        self.assertEqual((nolar(r), [(t["tabaka"], t["populasyon"], t["orneklem"]) for t in r["tabakalar"]], r["bilgi"]["agirlik"]), (nolar(r2), [("Yüksek", 5, 3), ("Orta", 10, 4), ("Düşük", 15, 3)], "Yüksek 3; Orta 2; Düşük 1"))
        self.assertTrue(all(k["neden"] == f"Risk tabakası: {k['risk']}" for k in r["orneklem"]))
        r = main.sec(pop(range(1, 9), ["Yüksek", "Belirsiz", "Belirsiz", "Düşük", "Düşük", "Düşük", "Düşük", "Orta"]), "risk", 4, 1, (D(10), D(1), D(0)))
        self.assertEqual(([(t["tabaka"], t["populasyon"], t["orneklem"]) for t in r["tabakalar"]], [u["tur"] for u in r["uyarilar"]]), ([("Yüksek", 1, 1), ("Orta", 1, 1), ("Düşük", 4, 0), ("Belirsiz", 2, 2)], ["Risk seviyesi belirsiz kalemler"]))


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "populasyon.csv", cikti, "rastgele", 10, 2026)
            self.assertEqual((r["ozet"]["populasyon"], r["ozet"]["orneklem"], r["ozet"]["populasyon_tutar"], r["ozet"]["kapsama"], r["uyarilar"]), (30, 10, D(1708600), D("33.33"), []))
            self.assertEqual([(t["tabaka"], t["populasyon"], t["populasyon_tutar"]) for t in r["tabakalar"]], [("Yüksek", 5, D(1210000)), ("Orta", 10, D(370000)), ("Düşük", 15, D(128600))])
            self.assertEqual(nolar(r), nolar(main.calistir(ORNEK / "populasyon.csv", cikti, "rastgele", 10, 2026)))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Örneklem", "Özet", "Uyarılar"])
            r = main.calistir(ORNEK / "populasyon.csv", cikti, "tutar", 10, 2026)
            self.assertEqual((sorted(k["no"] for k in r["orneklem"] if k["neden"].startswith("Kesin")), len(r["orneklem"]), r["bilgi"]["aralik"]), (["IS-003", "IS-009", "IS-019", "IS-027"], 10, D("98933.33")))
            self.assertTrue(r["ozet"]["kapsama_tutar"] > D(65))
            r = main.calistir(ORNEK / "populasyon.csv", cikti, "risk", 10, 2026)
            self.assertEqual([(t["tabaka"], t["orneklem"]) for t in r["tabakalar"]], [("Yüksek", 3), ("Orta", 4), ("Düşük", 3)])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "p.csv"
            y.write_text("Fiş No;İşlem Tarihi;İşlem Tutarı;Risk Düzeyi;İşlem Türü\nF1;05.09.2026;1.250,50;YÜKSEK;EFT\nF2;;;;\nf1;;5;;\nF3;;x;;\n;;5;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual((l, [(u["tur"], u["kim"]) for u in uy]),
                             ([{"no": "F1", "tarih": date(2026, 9, 5), "tutar": D("1250.50"), "risk": "Yüksek", "aciklama": "EFT"}, {"no": "F2", "tarih": None, "tutar": None, "risk": "Belirsiz", "aciklama": ""}],
                              [("Mükerrer işlem numarası", "f1"), ("Satır okunamadı", "F3")]))
            self.assertEqual([main.agirlik_coz(x) for x in ("3;2;1", "4 2 1", "1;0;0", "0;0;0", "3;2", "a;b;c", "3;2;-1")], [(D(3), D(2), D(1)), (D(4), D(2), D(1)), (D(1), D(0), D(0)), None, None, None, None])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--populasyon", str(y), "--cikti", c]
            self.assertEqual([main.main(arg + v) for v in ([], ["--yontem", "tutar", "--adet", "1", "--tohum", "7"], ["--yontem", "risk", "--adet", "2", "--agirlik", "4;2;1"])], [0, 0, 0])
            self.assertEqual([main.main(arg + v) for v in (["--adet", "0"], ["--agirlik", "1;2"], ["--agirlik", "0;0;0"])], [2] * 3)
            self.assertEqual(main.main(["--populasyon", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Tarih;Tutar\n05.09.2026;1\n", "İşlem No;Tutar\n;1\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            y.write_text("İşlem No;Tutar\nA;0\n", encoding="utf-8")
            self.assertEqual(main.main(arg + ["--yontem", "tutar"]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
