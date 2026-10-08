import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
SIMDI = datetime(2026, 10, 8, 9, 0)


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "s.xlsx"
        cls.s = main.calistir(ORNEK / "seferler.csv", ORNEK / "asama_sureleri.csv", cls.cikti, SIMDI)
        cls.x = {x.no: x for x in cls.s["seferler"]}
        cls.u = {}
        for u in cls.s["uyarilar"]:
            cls.u.setdefault(u["tur"], []).append(u)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_tahmin(self):
        self.assertEqual({k: (x.durum, x.bitis) for k, x in self.x.items()},
                         {"S-2401": ("Teslim edildi", datetime(2026, 10, 3, 22)), "S-2402": ("Yolda", datetime(2026, 10, 9, 16)),
                          "S-2403": ("Gecikecek", datetime(2026, 10, 11, 3)), "S-2404": ("Yolda", datetime(2026, 10, 9, 4)),
                          "S-2405": ("Yolda", datetime(2026, 10, 9, 5)), "S-2406": ("Planlandı", datetime(2026, 10, 11, 21)),
                          "S-2407": ("Planlandı", datetime(2026, 10, 11, 22)), "S-2408": ("Planlandı", datetime(2026, 10, 12, 20))})
        self.assertEqual(self.x["S-2403"].tahmini["Sınır Çıkış"], SIMDI)
        self.assertEqual(self.x["S-2403"].fark_saat, 57)
        self.assertEqual(self.x["S-2401"].fark_saat, -44)

    def test_uyarilar(self):
        self.assertEqual([u["tur"] for u in self.s["uyarilar"] if u["onem"] == "Yüksek"], ["Teslim gecikecek", "Yükleme gecikti", "Araç çakışması"])
        self.assertIn("2 gün 9 saat geç", self.u["Teslim gecikecek"][0]["aciklama"])
        self.assertIn("1 gün 1 saat geçti", self.u["Yükleme gecikti"][0]["aciklama"])
        self.assertEqual(self.u["Araç çakışması"][0]["sefer"], "S-2408")
        gecikme = {u["sefer"]: u["aciklama"] for u in self.u["Aşama gecikti"]}
        self.assertIn("Sınır Çıkış (Habur)", gecikme["S-2405"])
        self.assertIn("21 saat geçti", gecikme["S-2405"])
        self.assertEqual([u["sefer"] for u in self.u["Konum bilgisi eski"]], ["S-2403"])
        self.assertEqual([u["sefer"] for u in self.u["CMR yok"]], ["S-2404"])

    def test_arac_durumu(self):
        a = {x["plaka"]: x for x in self.s["araclar"]}
        self.assertEqual((a["34 ABC 101"]["aktif"], a["34 ABC 101"]["musait"]), (None, datetime(2026, 10, 3, 22)))
        self.assertEqual((a["34 ABC 102"]["aktif"].no, a["34 ABC 102"]["musait"]), ("S-2402", datetime(2026, 10, 9, 16)))
        self.assertEqual([y.no for y in a["34 ABC 102"]["sirada"]], ["S-2408"])

    def test_excel(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Seferler", "Araç Durumu", "Aşama Matrisi", "Performans", "Uyarılar"])
        self.assertEqual(wb["Seferler"]["A2"].value, "S-2403")
        self.assertEqual(wb["Seferler"]["P1"].value, "Operasyon Notu")


class KuralTesti(unittest.TestCase):
    def test_tanimsiz_ulke_varsayilan_ve_sira_hatasi(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t / "d.csv").write_text("Varış Ülke;Aşama;Süre (saat)\n*;Teslim;10\n", encoding="utf-8")
            (t / "s.csv").write_text("Sefer No;Çekici Plaka;Varış Ülke;Planlanan Teslim;Yükleme;Teslim\n"
                                     "A;06 X 1;İtalya;08.10.2026 20:00;08.10.2026 08:00;\n"
                                     "B;06 X 2;İtalya;;08.10.2026 08:00;07.10.2026 08:00\n", encoding="utf-8")
            s = main.calistir(t / "s.csv", t / "d.csv", t / "o.xlsx", SIMDI)
            x = {y.no: y for y in s["seferler"]}
            self.assertEqual(x["A"].tahmini_teslim, datetime(2026, 10, 8, 18))
            turler = {(u["tur"], u["sefer"]) for u in s["uyarilar"]}
            self.assertIn(("Veri hatası", "B"), turler)
            self.assertIn(("Konum bilgisi eski", "A"), turler)

    def test_sure_tanimi_yok(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t / "d.csv").write_text("Varış Ülke;Aşama;Süre (saat)\nAlmanya;Teslim;70\n", encoding="utf-8")
            (t / "s.csv").write_text("Sefer No;Varış Ülke;Planlanan Yükleme\nA;Irak;09.10.2026 08:00\n", encoding="utf-8")
            s = main.calistir(t / "s.csv", t / "d.csv", t / "o.xlsx", SIMDI)
            self.assertEqual(s["seferler"][0].durum, "Süre tanımı yok")

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--seferler", str(Path(tmp) / "yok.csv")]), 1)
            self.assertEqual(main.main(["--simdi", "dün", "--cikti", str(Path(tmp) / "x.xlsx")]), 1)


if __name__ == "__main__":
    unittest.main()
