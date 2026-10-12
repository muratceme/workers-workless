"""
Final Kontrol Raporu — Workers / Workless kod bloğu
Kalite Kontrol › Kalite Kontrol Şefi

Final denetim kayıtlarından kritik, majör ve minör hata sayılarını AQL kabul sayılarıyla karşılaştırır; sonuç raporu üretir.
  Hata sayısı (sınıf) = denetimin o sınıftaki hata adetlerinin toplamı
  Sonuç               = herhangi bir sınıfta hata sayısı > kabul sayısı ise RET; hepsi ≤ kabul sayısı ise KABUL
                        (kritik kabul sayısı boşsa 0 alınır; majör veya minör kabul sayısı boşsa "Değerlendirilemedi")
  Hata oranı %        = hata sayısı ÷ numune adedi × 100
Kabul sayıları dosyadan okunur: örneklem planınızdaki (parti büyüklüğü, muayene seviyesi, AQL) değerleri siz girersiniz.
Plan üretmek için "AQL Örneklem Planı" paketini kullanabilirsiniz. Kod hatanın sınıfını değiştirmez; sınıflandırma denetçinindir.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


DEN_SUTUN = {"no": ("denetim no", "kontrol no", "rapor no"), "siparis": ("siparis no", "siparis"), "model": ("model", "urun"), "atolye": ("atolye", "uretici", "fason", "hat"),
             "musteri": ("musteri",), "parti": ("parti adedi", "parti buyuklugu", "parti"), "numune": ("numune adedi", "orneklem", "numune"),
             "ac_kritik": ("kabul kritik", "kritik kabul", "ac kritik"), "ac_major": ("kabul major", "major kabul", "ac major"), "ac_minor": ("kabul minor", "minor kabul", "ac minor"),
             "tarih": ("tarih", "denetim tarihi"), "denetci": ("denetci", "kontrolor")}
HATA_SUTUN = {"no": ("denetim no", "kontrol no", "rapor no"), "hata": ("hata", "hata turu", "hata tanimi"), "sinif": ("sinif", "hata sinifi", "derece"), "adet": ("adet", "hata adedi"),
              "bolge": ("bolge", "yer", "operasyon")}
SINIFLAR = (("kritik", "Kritik"), ("major", "Majör"), ("minor", "Minör"))


def sinif_bul(x) -> str | None:
    k = katla(x)
    return "kritik" if k.startswith(("kritik", "critical")) else "major" if k.startswith("maj") else "minor" if k.startswith("min") else None


def oku(den_yol: Path, hata_yol: Path) -> tuple[dict, list[dict], list[dict]]:
    uy, den, hatalar = [], {}, []
    for r in kayitlar(den_yol, DEN_SUTUN, ("no", "numune")):
        no = metin(r["no"])
        if not no:
            continue
        n, parti = para(r["numune"]), para(r["parti"])
        if no in den:
            uy.append(uyari("Yüksek", "Mükerrer denetim", no, f"{den_yol.name} satır {r['_satir']}: denetim numarası tekrar ediyor; ilk satır kullanıldı."))
            continue
        if n is None or n <= 0:
            uy.append(uyari("Yüksek", "Numune adedi okunamadı", no, f"{den_yol.name} satır {r['_satir']}: '{metin(r['numune'])}'. Denetim rapora alınmadı."))
            continue
        if parti is not None and n > parti:
            uy.append(uyari("Orta", "Numune partiden büyük", no, f"Numune {n:f} > parti {parti:f}. Adetleri kontrol edin."))
        ac = {}
        for k, ad in SINIFLAR:
            v = para(r[f"ac_{k}"])
            if metin(r[f"ac_{k}"]) and (v is None or v < 0 or v != v.to_integral_value()):
                uy.append(uyari("Yüksek", "Kabul sayısı okunamadı", no, f"{ad}: '{metin(r[f'ac_{k}'])}'. Tam sayı olmalı; boş sayıldı."))
                v = None
            ac[k] = int(v) if v is not None else (0 if k == "kritik" else None)
        den[no] = {"no": no, "siparis": metin(r["siparis"]), "model": metin(r["model"]), "atolye": metin(r["atolye"]) or "Belirtilmemiş", "musteri": metin(r["musteri"]), "parti": parti,
                   "numune": n, "ac": ac, "tarih": tarih(r["tarih"]), "denetci": metin(r["denetci"])}
    for r in kayitlar(hata_yol, HATA_SUTUN, ("no", "sinif")):
        no = metin(r["no"])
        if not no:
            continue
        s, a = sinif_bul(r["sinif"]), para(r["adet"]) if metin(r["adet"]) else Decimal(1)
        kim = f"{no} · {metin(r['hata'])}"
        if no not in den:
            uy.append(uyari("Yüksek", "Hatanın denetimi yok", kim, f"{hata_yol.name} satır {r['_satir']}: denetim listesinde bu numara yok; hata rapora alınmadı."))
        elif not s:
            uy.append(uyari("Yüksek", "Hata sınıfı tanınmadı", kim, f"{hata_yol.name} satır {r['_satir']}: '{metin(r['sinif'])}'. Kritik, Majör veya Minör yazın; hata rapora alınmadı."))
        elif a is None or a <= 0:
            uy.append(uyari("Yüksek", "Hata adedi okunamadı", kim, f"{hata_yol.name} satır {r['_satir']}: '{metin(r['adet'])}'. Hata rapora alınmadı."))
        else:
            hatalar.append({"no": no, "hata": metin(r["hata"]) or "Belirtilmemiş", "sinif": s, "adet": int(a), "bolge": metin(r["bolge"])})
    return den, hatalar, uy


def degerlendir(den: dict, hatalar: list[dict]) -> dict:
    if not den:
        raise ValueError("Değerlendirilecek denetim yok.")
    for d in den.values():
        d["hata"] = {k: sum(h["adet"] for h in hatalar if h["no"] == d["no"] and h["sinif"] == k) for k, _ in SINIFLAR}
        d["asan"] = [ad for k, ad in SINIFLAR if d["ac"][k] is not None and d["hata"][k] > d["ac"][k]]
        d["eksik"] = [ad for k, ad in SINIFLAR if d["ac"][k] is None]
        d["sonuc"] = "RET" if d["asan"] else "Değerlendirilemedi" if d["eksik"] else "KABUL"
        d["oran"] = {k: yuvarla(Decimal(d["hata"][k]) / d["numune"] * 100) for k, _ in SINIFLAR}
        d["gerekce"] = ("; ".join(f"{ad} hata {d['hata'][k]} > kabul {d['ac'][k]}" for k, ad in SINIFLAR if ad in d["asan"])
                        or (f"{', '.join(d['eksik'])} kabul sayısı girilmemiş" if d["eksik"] else "Tüm sınıflar kabul sayısı içinde"))

    def pareto(l):
        g = defaultdict(int)
        for h in l:
            g[(h["hata"], h["sinif"])] += h["adet"]
        top, bir, out = sum(g.values()), 0, []
        for (hata, sinif), a in sorted(g.items(), key=lambda x: (-x[1], x[0])):
            bir += a
            out.append({"hata": hata, "sinif": dict(SINIFLAR)[sinif], "adet": a, "pay": yuvarla(Decimal(a) / top * 100), "birikimli": yuvarla(Decimal(bir) / top * 100)})
        return out

    atolye = {}
    for d in den.values():
        a = atolye.setdefault(d["atolye"], {"atolye": d["atolye"], "denetim": 0, "kabul": 0, "ret": 0, "belirsiz": 0, "numune": Decimal(0), "kritik": 0, "major": 0, "minor": 0, "nolar": set()})
        a["denetim"] += 1
        a["kabul" if d["sonuc"] == "KABUL" else "ret" if d["sonuc"] == "RET" else "belirsiz"] += 1
        a["numune"] += d["numune"]
        a["nolar"].add(d["no"])
        for k, _ in SINIFLAR:
            a[k] += d["hata"][k]
    for a in atolye.values():
        a["ret_oran"] = yuvarla(Decimal(a["ret"]) / a["denetim"] * 100)
        a["major_oran"] = yuvarla(Decimal(a["major"]) / a["numune"] * 100)
        nolar = a.pop("nolar")
        p = pareto([h for h in hatalar if h["no"] in nolar])
        a["en_sik"] = f"{p[0]['hata']} ({p[0]['adet']})" if p else ""
    dl = sorted(den.values(), key=lambda d: (("RET", "Değerlendirilemedi", "KABUL").index(d["sonuc"]), d["no"]))
    return {"denetimler": dl, "pareto": pareto(hatalar), "atolye": sorted(atolye.values(), key=lambda a: (-a["ret_oran"], a["atolye"])),
            "say": {s: sum(d["sonuc"] == s for d in dl) for s in ("KABUL", "RET", "Değerlendirilemedi")}, "hatalar": hatalar}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    renk = {"KABUL": YESIL, "RET": KIRMIZI, "Değerlendirilemedi": SARI}
    ws = wb.active
    ws.title = "Denetim Sonuçları"
    _baslik(ws, ["Denetim No", "Tarih", "Sipariş No", "Model", "Atölye", "Müşteri", "Parti Adedi", "Numune Adedi", "Kritik (hata / kabul)", "Majör (hata / kabul)", "Minör (hata / kabul)",
                 "Majör Hata %", "Sonuç", "Gerekçe", "Denetçi", "Aksiyon (siz doldurun)"], (12, 12, 12, 16, 20, 18, 10, 10, 12, 12, 12, 10, 18, 50, 16, 30))
    for d in r["denetimler"]:
        hk = lambda k: f"{d['hata'][k]} / {d['ac'][k] if d['ac'][k] is not None else '—'}"   # noqa: E731
        _satir(ws, [d["no"], d["tarih"], d["siparis"], d["model"], d["atolye"], d["musteri"], d["parti"], d["numune"], hk("kritik"), hk("major"), hk("minor"), d["oran"]["major"], d["sonuc"],
                    d["gerekce"], d["denetci"], ""], para_s=(12,), tarih_s=(2,), dolgu={13: renk[d["sonuc"]]}, sar=(14,))
    ws = wb.create_sheet("Hata Dağılımı")
    _baslik(ws, ["Hata", "Sınıf", "Adet", "Pay %", "Birikimli %"], (36, 10, 9, 10, 12))
    for p in r["pareto"]:
        _satir(ws, [p["hata"], p["sinif"], p["adet"], p["pay"], p["birikimli"]], para_s=(4, 5), dolgu={2: KIRMIZI if p["sinif"] == "Kritik" else None})
    ws = wb.create_sheet("Atölye Bazında")
    _baslik(ws, ["Atölye", "Denetim", "Kabul", "Ret", "Değerlendirilemedi", "Ret %", "Numune Toplamı", "Kritik Hata", "Majör Hata", "Minör Hata", "Majör Hata %", "En Sık Hata"],
            (24, 9, 8, 8, 14, 9, 12, 10, 10, 10, 11, 30))
    for a in r["atolye"]:
        _satir(ws, [a["atolye"], a["denetim"], a["kabul"], a["ret"] or None, a["belirsiz"] or None, a["ret_oran"], a["numune"], a["kritik"] or None, a["major"], a["minor"], a["major_oran"], a["en_sik"]],
               para_s=(6, 11), dolgu={4: KIRMIZI if a["ret"] else None, 8: KIRMIZI if a["kritik"] else None})
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(den_yol: Path, hata_yol: Path, cikti: Path) -> dict:
    den, hatalar, uy = oku(den_yol, hata_yol)
    r = degerlendir(den, hatalar)
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Final denetim hata sayılarını AQL kabul sayılarıyla karşılaştırıp sonuç raporu üretir.")
    p.add_argument("--denetimler", type=Path, default=ORNEK / "denetimler.csv",
                   help="Denetim No, Sipariş No, Model, Atölye, Müşteri, Parti Adedi, Numune Adedi, Kabul Kritik, Kabul Majör, Kabul Minör, Tarih, Denetçi")
    p.add_argument("--hatalar", type=Path, default=ORNEK / "hatalar.csv", help="Denetim No, Hata, Sınıf (Kritik / Majör / Minör), Adet, Bölge")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "final_kontrol_raporu.xlsx")
    a = p.parse_args(argv)
    for y in (a.denetimler, a.hatalar):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.denetimler, a.hatalar, a.cikti)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    s = r["say"]
    print(f"[OK] {len(r['denetimler'])} denetim · {s['KABUL']} kabul · {s['RET']} ret · {s['Değerlendirilemedi']} değerlendirilemedi")
    for d in r["denetimler"]:
        if d["sonuc"] != "KABUL":
            print(f"[!] {d['no']} · {d['atolye']} · {d['sonuc']}: {d['gerekce']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
