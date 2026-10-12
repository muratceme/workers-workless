import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def den(no="D1", numune=80, kritik=0, major=5, minor=7, atolye="A"):
    return {no: {"no": no, "siparis": "", "model": "", "atolye": atolye, "musteri": "", "parti": D(1200), "numune": D(numune), "ac": {"kritik": kritik, "major": major, "minor": minor},
                 "tarih": None, "denetci": ""}}


def h(sinif, adet, hata="Leke", no="D1"):
    return {"no": no, "hata": hata, "sinif": sinif, "adet": adet, "bolge": ""}


def tek(hatalar, **kw):
    return main.degerlendir(den(**kw), hatalar)["denetimler"][0]


class DegerlendirmeTesti(unittest.TestCase):
    def test_kabul_ret_sinirlari(self):
        self.assertEqual(tek([h("major", 5), h("minor", 7)])["sonuc"], "KABUL")                    # kabul sayısına eşit
        d = tek([h("major", 3), h("major", 3, "Dikiş")])
        self.assertEqual((d["sonuc"], d["hata"]["major"], d["asan"], d["gerekce"], d["oran"]["major"]), ("RET", 6, ["Majör"], "Majör hata 6 > kabul 5", D("7.50")))
        self.assertEqual(tek([h("minor", 8)])["sonuc"], "RET")
        self.assertEqual(tek([h("kritik", 1)])["gerekce"], "Kritik hata 1 > kabul 0")
        self.assertEqual(tek([])["sonuc"], "KABUL")

    def test_eksik_kabul_sayisi(self):
        d = tek([h("major", 2), h("minor", 30)], minor=None)
        self.assertEqual((d["sonuc"], d["gerekce"]), ("Değerlendirilemedi", "Minör kabul sayısı girilmemiş"))
        self.assertEqual(tek([h("major", 9)], minor=None)["sonuc"], "RET")                         # girilen sınıf aşılmışsa yine RET

    def test_pareto_ve_atolye(self):
        dn = den("D1") | den("D2", 20, atolye="B") | den("D3", 100, atolye="B")
        hl = [h("minor", 6, "İplik"), h("major", 2), h("major", 9, "Dikiş", "D2"), h("minor", 2, "İplik", "D2"), h("kritik", 1, "İğne", "D3")]
        r = main.degerlendir(dn, hl)
        self.assertEqual([(p["hata"], p["sinif"], p["adet"], p["pay"], p["birikimli"]) for p in r["pareto"]],
                         [("Dikiş", "Majör", 9, D("45.00"), D("45.00")), ("İplik", "Minör", 8, D("40.00"), D("85.00")), ("Leke", "Majör", 2, D("10.00"), D("95.00")), ("İğne", "Kritik", 1, D("5.00"), D("100.00"))])
        self.assertEqual([(a["atolye"], a["denetim"], a["ret"], a["ret_oran"], a["numune"], a["major"], a["major_oran"], a["en_sik"]) for a in r["atolye"]],
                         [("B", 2, 2, D("100.00"), D(120), 9, D("7.50"), "Dikiş (9)"), ("A", 1, 0, D("0.00"), D(80), 2, D("2.50"), "İplik (6)")])
        self.assertEqual(([d["no"] for d in r["denetimler"]], r["say"]), (["D2", "D3", "D1"], {"KABUL": 1, "RET": 2, "Değerlendirilemedi": 0}))
        with self.assertRaises(ValueError):
            main.degerlendir({}, [])
        self.assertEqual([main.sinif_bul(x) for x in ("KRİTİK", "Major", "minör", "Min.", "A")], ["kritik", "major", "minor", "minor", None])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "f.xlsx"
            r = main.calistir(ORNEK / "denetimler.csv", ORNEK / "hatalar.csv", cikti)
            self.assertEqual({d["no"]: (d["sonuc"], d["hata"]) for d in r["denetimler"]},
                             {"FD-01": ("KABUL", {"kritik": 0, "major": 4, "minor": 6}), "FD-02": ("RET", {"kritik": 0, "major": 5, "minor": 2}), "FD-03": ("RET", {"kritik": 1, "major": 3, "minor": 4}),
                              "FD-04": ("KABUL", {"kritik": 0, "major": 0, "minor": 0}), "FD-05": ("Değerlendirilemedi", {"kritik": 0, "major": 2, "minor": 9})})
            self.assertEqual([(p["hata"], p["adet"]) for p in r["pareto"]][:4], [("İplik ucu", 16), ("Leke", 6), ("Ütü izi", 5), ("Dikiş atlaması", 4)])
            self.assertEqual({a["atolye"][-1]: (a["denetim"], a["ret"], a["belirsiz"], a["major_oran"]) for a in r["atolye"]}, {"A": (2, 0, 0, D("4.00")), "B": (2, 2, 0, D("4.57")), "C": (1, 0, 1, D("2.50"))})
            self.assertEqual([u["tur"] for u in r["uyarilar"]], ["Hatanın denetimi yok"])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Denetim Sonuçları", "Hata Dağılımı", "Atölye Bazında", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            d, hy = Path(tmp) / "d.csv", Path(tmp) / "h.csv"
            d.write_text("Denetim No;Parti Adedi;Numune Adedi;Kabul Kritik;Kabul Majör;Kabul Minör\nD1;50;80;;2,5;7\nD1;;10;;;\nD2;;yok;;;\n", encoding="utf-8")
            hy.write_text("Denetim No;Hata;Sınıf;Adet\nD1;Leke;Majör;\nD1;Leke;Orta;1\nD1;Leke;Minör;sıfır\nD9;Leke;Minör;1\n", encoding="utf-8")
            dn, hl, uy = main.oku(d, hy)
            self.assertEqual((dn["D1"]["ac"], [(x["sinif"], x["adet"]) for x in hl]), ({"kritik": 0, "major": None, "minor": 7}, [("major", 1)]))
            self.assertEqual(sorted(u["tur"] for u in uy), ["Hata adedi okunamadı", "Hata sınıfı tanınmadı", "Hatanın denetimi yok", "Kabul sayısı okunamadı", "Mükerrer denetim",
                                                            "Numune adedi okunamadı", "Numune partiden büyük"])
            c = str(Path(tmp) / "x.xlsx")
            self.assertEqual(main.main(["--denetimler", str(d), "--hatalar", str(hy), "--cikti", c]), 0)
            self.assertEqual(main.main(["--denetimler", str(Path(tmp) / "yok.csv")]), 1)
            d.write_text("Denetim No;Numune Adedi\n", encoding="utf-8")
            self.assertEqual(main.main(["--denetimler", str(d), "--hatalar", str(hy), "--cikti", c]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
