"""
Yeniden Sipariş Noktası Hesabı — Workers / Workless kod bloğu
Üretim › Üretim Planlama / Malzeme Planlama

Her stok kalemi için emniyet stoğunu, yeniden sipariş noktasını (YSN) ve ekonomik sipariş miktarını (ESM)
hesaplar; stok pozisyonu (eldeki + yoldaki − ayrılmış) YSN'nin altındaysa sipariş önerir:

    Emniyet stoğu = z × √( L × σd² + d² × σL² )        (talep ve tedarik süresi belirsizliği birlikte)
    YSN           = d × L + Emniyet stoğu
    ESM           = √( 2 × D × S / (h × c) )           (D yıllık talep, S sipariş maliyeti, h yıllık stok tutma oranı, c birim maliyet)

d ve σd günlük talebin ortalaması ve standart sapmasıdır; z, hizmet düzeyinin (döngü başına stoksuz kalmama
olasılığı) standart normal karşılığıdır. Talep, tüketim geçmişinden (dönem toplamları) ya da doğrudan
verilen değerlerden alınır. Sipariş miktarı en az sipariş miktarına ve ambalaj katına yuvarlanır.
İnternete bağlanmaz.

Kullanım:
    python main.py                                           # örnek kalemlerle dener
    python main.py --girdi kalemler.xlsx --gecmis tuketim.xlsx --donem-gun 30
    python main.py --girdi kalemler.xlsx --hizmet 97,5 --siparis-maliyeti 750 --tutma-orani 25
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from statistics import NormalDist, fmean, pstdev

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def sayi(x, varsayilan: float | None = None) -> float | None:
    if x in (None, ""):
        return varsayilan
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip().replace("TL", "").replace("%", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(Decimal(s))
    except InvalidOperation:
        return varsayilan


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


ALANLAR = {
    "kod": ("stok kodu", "malzeme kodu", "ürün kodu", "kod", "sku"),
    "ad": ("stok adı", "malzeme adı", "ürün adı", "ad", "açıklama"),
    "d": ("ortalama günlük talep", "günlük talep", "günlük ortalama tüketim"),
    "sd": ("günlük talep std", "günlük talep standart sapma", "talep std"),
    "L": ("tedarik süresi (gün)", "tedarik süresi", "temin süresi", "lead time"),
    "sL": ("tedarik süresi std", "tedarik süresi standart sapma", "temin süresi std"),
    "hizmet": ("hizmet düzeyi", "hizmet düzeyi %"),
    "eldeki": ("eldeki stok", "mevcut stok", "stok"),
    "yoldaki": ("açık sipariş", "yoldaki", "yoldaki stok", "açık satın alma"),
    "ayrilmis": ("ayrılmış", "rezerve", "ayrılmış stok"),
    "maliyet": ("birim maliyet", "birim fiyat"),
    "moq": ("en az sipariş", "min. sipariş", "moq", "minimum sipariş"),
    "kat": ("ambalaj katı", "koli içi", "sipariş katı"),
    "birim": ("birim",),
}


def sutunlar(baslik):
    b = [kucuk(x) for x in baslik]
    return {alan: next((i for i, x in enumerate(b) if x in adlar), None) for alan, adlar in ALANLAR.items()}


def donem_anahtari(x) -> str | None:
    if isinstance(x, (datetime, date)):
        return x.strftime("%Y-%m-%d")
    s = str(x or "").strip()
    m = re.match(r"^(\d{4})[-/.](\d{1,2})(?:[-/.](\d{1,2}))?", s)
    if m:
        return f"{m.group(1)}-{int(m.group(2)):02d}" + (f"-{int(m.group(3)):02d}" if m.group(3) else "")
    m = re.match(r"^(\d{1,2})[./](\d{1,2})[./](\d{4})$", s)
    if m:
        return f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
    m = re.match(r"^(\d{1,2})[./-](\d{4})$", s)
    if m:
        return f"{m.group(2)}-{int(m.group(1)):02d}"
    return None


def gecmis_oku(yol: Path | None, donem_gun: float) -> dict[str, tuple[float, float, int]]:
    """Tüketim geçmişinden günlük ortalama ve standart sapma: dönem toplamları → d = ort/gün, σd = σ/√gün
    (günler arası talebin bağımsız olduğu varsayımı)."""
    if not yol:
        return {}
    s = tablo_oku(yol)
    k = sutunlar(s[0])
    b = [kucuk(x) for x in s[0]]
    i_t = next((i for i, x in enumerate(b) if x in ("tarih", "dönem", "ay", "hafta")), None)
    i_m = next((i for i, x in enumerate(b) if x in ("miktar", "tüketim", "çıkış miktarı", "satış miktarı")), None)
    if k["kod"] is None or i_t is None or i_m is None:
        raise SystemExit(f"Geçmiş dosyasında Stok Kodu, Tarih/Dönem ve Miktar gerekli. Başlıklar: {s[0]}")
    toplam: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    donemler = set()
    for r in s[1:]:
        d = donem_anahtari(r[i_t])
        if d and r[k["kod"]]:
            toplam[str(r[k["kod"]]).strip()][d] += sayi(r[i_m], 0.0)
            donemler.add(d)
    donemler = sorted(donemler)
    sonuc = {}
    for kod, v in toplam.items():
        seri = [v.get(p, 0.0) for p in donemler]
        ilk = next((i for i, x in enumerate(seri) if x), len(seri))
        seri = seri[ilk:]                                 # yeni kalemde ilk tüketimden öncesi sayılmaz
        if seri:
            sonuc[kod] = (fmean(seri) / donem_gun, pstdev(seri) / math.sqrt(donem_gun), len(seri))
    return sonuc


# ----------------------------------------------------------------------------
# Hesap
# ----------------------------------------------------------------------------

def emniyet_stogu(z: float, d: float, sd: float, L: float, sL: float) -> float:
    return z * math.sqrt(L * sd ** 2 + d ** 2 * sL ** 2)


def esm(yillik_talep: float, siparis_maliyeti: float, tutma_orani: float, birim_maliyet: float) -> float | None:
    h = tutma_orani * birim_maliyet
    if yillik_talep <= 0 or siparis_maliyeti <= 0 or h <= 0:
        return None
    return math.sqrt(2 * yillik_talep * siparis_maliyeti / h)


def yuvarla_kat(miktar: float, moq: float, kat: float) -> float:
    m = max(miktar, moq or 0)
    if kat and kat > 0:
        m = math.ceil(m / kat - 1e-9) * kat
    return math.ceil(m - 1e-9) if not kat else m


def calistir(girdi: Path, cikti: Path, gecmis_yolu: Path | None = None, donem_gun: float = 30, hizmet: float = 95,
             siparis_maliyeti: float = 0, tutma_orani: float = 0.25, yil_gun: float = 365) -> dict:
    s = tablo_oku(girdi)
    k = sutunlar(s[0])
    if k["kod"] is None or k["L"] is None:
        raise SystemExit(f"Kalem listesinde Stok Kodu ve Tedarik Süresi (gün) gerekli. Başlıklar: {s[0]}")
    gecmis = gecmis_oku(gecmis_yolu, donem_gun)
    al = lambda r, a: r[k[a]] if k[a] is not None and k[a] < len(r) else None  # noqa: E731
    kalemler, uyarilar = [], []
    for r in s[1:]:
        kod = str(al(r, "kod") or "").strip()
        if not kod:
            continue
        d, sd, kaynak = sayi(al(r, "d")), sayi(al(r, "sd")), "dosya"
        if (d is None or sd is None) and kod in gecmis:
            gd, gsd, n = gecmis[kod]
            d = gd if d is None else d
            sd = gsd if sd is None else sd
            kaynak = f"geçmiş ({n} dönem)"
            if n < 6:
                uyarilar.append(f"{kod}: talep yalnız {n} dönemden hesaplandı; emniyet stoğu güvenilir olmayabilir")
        if d is None:
            uyarilar.append(f"{kod}: günlük talep yok (dosyada veya geçmişte); hesaplanmadı")
            continue
        sd = sd or 0.0
        L, sL = sayi(al(r, "L"), 0.0), sayi(al(r, "sL"), 0.0)
        hz = sayi(al(r, "hizmet"), hizmet)
        hz = hz / 100 if hz > 1 else hz
        if not 0.5 <= hz < 1:
            raise SystemExit(f"{kod}: hizmet düzeyi %50 ile %100 (hariç) arasında olmalı")
        z = NormalDist().inv_cdf(hz)
        ss = emniyet_stogu(z, d, sd, L, sL)
        ysn = d * L + ss
        c = sayi(al(r, "maliyet"), 0.0)
        e = esm(d * yil_gun, siparis_maliyeti, tutma_orani, c)
        eldeki, yoldaki, ayrilmis = sayi(al(r, "eldeki"), 0.0), sayi(al(r, "yoldaki"), 0.0), sayi(al(r, "ayrilmis"), 0.0)
        pozisyon = eldeki + yoldaki - ayrilmis
        moq, kat = sayi(al(r, "moq"), 0.0), sayi(al(r, "kat"), 0.0)
        oneri = 0.0
        if pozisyon <= ysn:
            # ESM varsa ESM; yoksa pozisyonu YSN + bir tedarik süresi talebine tamamlayacak miktar
            hedef = e if e else ysn + d * L - pozisyon
            oneri = yuvarla_kat(max(hedef, ysn - pozisyon), moq, kat)
        gun_kaldi = (eldeki - ayrilmis) / d if d > 0 else None
        kalemler.append({"kod": kod, "ad": str(al(r, "ad") or ""), "birim": str(al(r, "birim") or ""), "d": d, "sd": sd, "kaynak": kaynak,
                         "L": L, "sL": sL, "hizmet": hz, "z": z, "ss": ss, "ysn": ysn, "esm": e, "maliyet": c, "eldeki": eldeki,
                         "yoldaki": yoldaki, "ayrilmis": ayrilmis, "pozisyon": pozisyon, "oneri": oneri, "gun_kaldi": gun_kaldi,
                         "acil": gun_kaldi is not None and gun_kaldi < L and pozisyon <= ysn,
                         "ss_deger": ss * c})
    _rapor(kalemler, cikti, uyarilar, siparis_maliyeti, tutma_orani, donem_gun)
    return {"kalemler": kalemler, "uyarilar": uyarilar}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
SARI = PatternFill("solid", fgColor="FFF4CE")
PARA = "#,##0.00"
MIKTAR = "#,##0.0"


def _rapor(kalemler, cikti, uyarilar, siparis_maliyeti, tutma_orani, donem_gun):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sipariş Önerisi"
    bas = ["Stok Kodu", "Stok Adı", "Birim", "Günlük Talep", "Günlük Talep σ", "Talep Kaynağı", "Tedarik Süresi (gün)", "Tedarik σ (gün)",
           "Hizmet Düzeyi", "z", "Emniyet Stoğu", "YSN", "ESM", "Eldeki", "Yoldaki", "Ayrılmış", "Stok Pozisyonu", "Önerilen Sipariş",
           "Eldeki Kaç Gün Yeter", "Durum", "Emniyet Stoğu Değeri (TL)"]
    ws.append(bas)
    for h in ws[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for x in sorted(kalemler, key=lambda x: (not x["acil"], x["oneri"] == 0, x["kod"])):
        durum = "ACİL: tedarik süresinden önce tükenir" if x["acil"] else "Sipariş ver" if x["oneri"] else "Yeterli"
        ws.append([x["kod"], x["ad"], x["birim"], x["d"], x["sd"], x["kaynak"], x["L"], x["sL"], x["hizmet"], round(x["z"], 3),
                   x["ss"], x["ysn"], x["esm"], x["eldeki"], x["yoldaki"], x["ayrilmis"], x["pozisyon"], x["oneri"] or None,
                   x["gun_kaldi"], durum, x["ss_deger"]])
        r = ws.max_row
        for c in (4, 5, 11, 12, 13, 14, 15, 16, 17, 18, 19):
            ws.cell(r, c).number_format = MIKTAR
        ws.cell(r, 9).number_format = "0.0%"
        ws.cell(r, 21).number_format = PARA
        if x["acil"]:
            ws.cell(r, 20).fill = KIRMIZI
        elif x["oneri"]:
            ws.cell(r, 20).fill = SARI
    for j, w in enumerate((12, 28, 7, 12, 12, 16, 11, 10, 9, 7, 12, 11, 11, 10, 10, 10, 12, 13, 12, 34, 15), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    if uyarilar:
        ws.append([])
        for u in uyarilar:
            ws.append(["Uyarı", u])

    b = wb.create_sheet("Bilgi")
    for s in [["Emniyet stoğu", "z × √(L × σd² + d² × σL²) — talep ve tedarik süresi belirsizliği birlikte; σL = 0 ise z × σd × √L"],
              ["YSN", "d × L + emniyet stoğu: stok pozisyonu bu seviyeye inince sipariş verilir"],
              ["ESM", f"√(2 × yıllık talep × sipariş maliyeti / (stok tutma oranı × birim maliyet)); sipariş maliyeti "
                      f"{siparis_maliyeti:g} TL, tutma oranı %{tutma_orani * 100:g}" + (" — sipariş maliyeti verilmediği için ESM hesaplanmadı"
                                                                                       if not siparis_maliyeti else "")],
              ["Stok pozisyonu", "Eldeki + yoldaki (açık sipariş) − ayrılmış (rezerve)"],
              ["Önerilen sipariş", "Pozisyon ≤ YSN ise: ESM (yoksa YSN + L günlük talep − pozisyon), en az YSN − pozisyon; en az sipariş "
                                   "miktarına ve ambalaj katına yukarı yuvarlanır"],
              ["Talep geçmişten", f"Dönem toplamları ({donem_gun:g} gün) → d = ortalama / {donem_gun:g}, σd = σ / √{donem_gun:g} "
                                  "(günlük talepler bağımsız varsayılır)"],
              ["Hizmet düzeyi", "Sipariş döngüsü başına stoksuz kalmama olasılığı (döngü hizmet düzeyi), z = NORM.S.TERS(hizmet)"],
              ["Not", "Talep normal dağılıma yakın olmayan (çok düzensiz, Z sınıfı) kalemlerde formül emniyet stoğunu hatalı verebilir"]]:
        b.append(s)
    b.column_dimensions["A"].width = 18
    b.column_dimensions["B"].width = 130
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description="Emniyet stoğu, yeniden sipariş noktası ve sipariş önerisi hesaplar.")
    ap.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "kalemler.csv",
                    help="Kalem listesi (.xlsx/.csv): Stok Kodu, Tedarik Süresi (gün), [Ortalama Günlük Talep, Günlük Talep Std, "
                         "Tedarik Süresi Std, Hizmet Düzeyi, Eldeki Stok, Açık Sipariş, Ayrılmış, Birim Maliyet, En Az Sipariş, Ambalaj Katı]")
    ap.add_argument("--gecmis", type=Path, help="Tüketim geçmişi (.xlsx/.csv: Stok Kodu, Tarih/Dönem, Miktar) — talep listede yoksa")
    ap.add_argument("--donem-gun", type=float, default=30, help="Geçmişteki bir dönemin gün sayısı (aylık 30, haftalık 7)")
    ap.add_argument("--hizmet", type=str, default="95", help="Varsayılan hizmet düzeyi, %% (kalemde verilmezse)")
    ap.add_argument("--siparis-maliyeti", type=str, default="0", help="Sipariş başına sabit maliyet, TL (ESM için)")
    ap.add_argument("--tutma-orani", type=str, default="25", help="Yıllık stok tutma maliyeti oranı, %% (ESM için)")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "yeniden_siparis.xlsx")
    a = ap.parse_args(argv)
    if a.girdi == BURASI / "ornek_veri" / "kalemler.csv":
        a.gecmis = a.gecmis or BURASI / "ornek_veri" / "aylik_tuketim.csv"
        a.siparis_maliyeti = a.siparis_maliyeti if a.siparis_maliyeti != "0" else "750"
    tutma = sayi(a.tutma_orani, 25)
    s = calistir(a.girdi, a.cikti, a.gecmis, a.donem_gun, sayi(a.hizmet, 95), sayi(a.siparis_maliyeti, 0),
                 tutma / 100 if tutma > 1 else tutma)
    acil = [x for x in s["kalemler"] if x["acil"]]
    oneri = [x for x in s["kalemler"] if x["oneri"]]
    print(f"[OK] {len(s['kalemler'])} kalem · sipariş önerisi: {len(oneri)} · acil: {len(acil)}")
    for x in acil:
        print(f"[!] ACİL {x['kod']} {x['ad'][:30]}: eldeki {x['gun_kaldi']:.1f} gün yeter, tedarik {x['L']:g} gün → {x['oneri']:g} {x['birim']}")
    for u in s["uyarilar"]:
        print(f"[!] {u}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
