"""
Satış Hedef-Gerçekleşme Raporu — Workers / Workless kod bloğu
Satış › Satış Destek Uzmanı

Satış hareketlerini aylık hedeflerle karşılaştırır:
  - Kırılımlar: temsilci, bölge, ürün grubu (hangi sütunlar varsa); ay, yıl başından bugüne (YTD) ve geçen yılın
    aynı dönemi (GY) karşılaştırması.
  - Ay içi rapor: geçen iş günü oranına göre beklenen gerçekleşme, gün sonu tahmini (run-rate) ve hedefe ulaşmak
    için kalan iş günlerinde gereken günlük satış.
  - İsteğe bağlı prim kademeleri (ör. %90 → 0,5 kat, %100 → 1 kat, %110 → 1,5 kat).
İnternete bağlanmaz.

Kullanım:
    python main.py                                          # örnek verilerle dener
    python main.py --satislar satislar.xlsx --hedefler hedefler.xlsx --ay 2026-09
    python main.py --satislar satislar.xlsx --hedefler hedefler.xlsx --ay 2026-10 --bugun 15.10.2026 --prim 90=0.5,100=1,110=1.5
"""
from __future__ import annotations

import argparse
import calendar
import csv
import re
import sys
from collections import OrderedDict, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
BOYUTLAR = OrderedDict([("temsilci", ("temsilci", "satış temsilcisi", "plasiyer", "satıcı")),
                        ("bolge", ("bölge", "region", "şube")),
                        ("grup", ("ürün grubu", "kategori", "ürün kategorisi", "marka"))])
BOYUT_AD = {"temsilci": "Temsilci", "bolge": "Bölge", "grup": "Ürün Grubu"}


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def katla(s) -> str:
    s = kucuk(s)
    for a, b in zip("ıöüşğç", "iousgc"):
        s = s.replace(a, b)
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def sayi(x) -> float:
    if x in (None, ""):
        return 0.0
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip().replace("TL", "").replace("₺", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1:
        s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        return 0.0


def tl(x: float, b: int = 0) -> str:
    return f"{x:,.{b}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    for b in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%Y %H:%M"):
        try:
            return datetime.strptime(str(x or "").strip()[:16] if "%H" in b else str(x or "").strip()[:10], b).date()
        except ValueError:
            continue
    return None


AYLAR = ["ocak", "subat", "mart", "nisan", "mayis", "haziran", "temmuz", "agustos", "eylul", "ekim", "kasim", "aralik"]


def ay_anahtari(x) -> str | None:
    if isinstance(x, (datetime, date)):
        return f"{x.year}-{x.month:02d}"
    s = katla(x)
    m = re.match(r"^(\d{4}) (\d{1,2})$", s) or re.match(r"^(\d{4})(\d{2})$", s)
    if m:
        return f"{m.group(1)}-{int(m.group(2)):02d}"
    m = re.match(r"^(\d{1,2}) (\d{4})$", s)
    if m:
        return f"{m.group(2)}-{int(m.group(1)):02d}"
    for i, a in enumerate(AYLAR, 1):
        if s.startswith(a):
            y = re.search(r"\d{4}", s)
            return f"{y.group(0)}-{i:02d}" if y else None
    d = tarih(x)
    return f"{d.year}-{d.month:02d}" if d else None


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def _bul(baslik, *adlar):
    b = [katla(x) for x in baslik]
    return next((b.index(katla(a)) for a in adlar if katla(a) in b), None)


def _al(r, i):
    return r[i] if i is not None and i < len(r) else None


# ----------------------------------------------------------------------------
# Okuma
# ----------------------------------------------------------------------------

def satislar_oku(yol: Path) -> tuple[list[dict], list[str]]:
    s = tablo_oku(yol)
    b = s[0]
    i_t = _bul(b, "tarih", "fatura tarihi", "sipariş tarihi")
    i_tut = _bul(b, "net tutar", "tutar", "net satış", "satış tutarı", "ciro")
    i_tur = _bul(b, "tür", "fatura türü", "işlem türü")
    if i_t is None or i_tut is None:
        raise SystemExit(f"Satış dosyasında Tarih ve Net Tutar sütunları gerekli. Başlıklar: {b}")
    i_b = {k: _bul(b, *v) for k, v in BOYUTLAR.items()}
    boyutlar = [k for k, i in i_b.items() if i is not None]
    satirlar = []
    for r in s[1:]:
        t = tarih(_al(r, i_t))
        if not t:
            continue
        tutar = sayi(_al(r, i_tut))
        if katla(_al(r, i_tur)).startswith("iade") and tutar > 0:
            tutar = -tutar
        satirlar.append({"tarih": t, "tutar": tutar, **{k: str(_al(r, i_b[k]) or "-").strip() for k in boyutlar}})
    return satirlar, boyutlar


def hedefler_oku(yol: Path, boyutlar: list[str]) -> tuple[dict[tuple, float], list[str]]:
    """(ay, boyut değerleri…) → hedef. Uzun (Ay, Temsilci, Hedef) veya geniş (Temsilci, 2026-01, 2026-02 …) biçim."""
    s = tablo_oku(yol)
    b = s[0]
    i_b = {k: _bul(b, *BOYUTLAR[k]) for k in boyutlar}
    hb = [k for k, i in i_b.items() if i is not None]               # hedefin hangi boyutlarda verildiği
    i_ay, i_h = _bul(b, "ay", "dönem", "yıl ay"), _bul(b, "hedef", "hedef tutar", "kota")
    hedef: dict[tuple, float] = defaultdict(float)
    if i_ay is not None and i_h is not None:
        for r in s[1:]:
            ay = ay_anahtari(_al(r, i_ay))
            if ay:
                hedef[(ay,) + tuple(str(_al(r, i_b[k]) or "-").strip() for k in hb)] += sayi(_al(r, i_h))
        return dict(hedef), hb
    aylar = {i: ay_anahtari(x) for i, x in enumerate(b) if i not in i_b.values() and ay_anahtari(x)}
    if not aylar:
        raise SystemExit(f"Hedef dosyasında Ay ve Hedef sütunları ya da ay başlıklı sütunlar (2026-01 …) gerekli. Başlıklar: {b}")
    for r in s[1:]:
        anahtar = tuple(str(_al(r, i_b[k]) or "-").strip() for k in hb)
        for i, ay in aylar.items():
            hedef[(ay,) + anahtar] += sayi(_al(r, i))
    return dict(hedef), hb


# ----------------------------------------------------------------------------
# Hesap
# ----------------------------------------------------------------------------

def is_gunu_say(bas: date, son: date, calisma_gunu: int = 5) -> int:
    """[bas, son] arası iş günü (Pzt–Cum veya Pzt–Cmt)."""
    n, d = 0, bas
    while d <= son:
        if d.weekday() < calisma_gunu:
            n += 1
        d += timedelta(days=1)
    return n


def prim_carpani(oran: float, kademeler: list[tuple[float, float]]) -> float:
    c = 0.0
    for esik, k in sorted(kademeler):
        if oran * 100 >= esik:
            c = k
    return c


def hesapla(satislar, boyutlar, hedef, hb, ay: str, bugun: date, calisma_gunu: int, kademeler) -> dict:
    yil, aynum = int(ay[:4]), int(ay[5:])
    ay_bas = date(yil, aynum, 1)
    ay_son = date(yil, aynum, calendar.monthrange(yil, aynum)[1])
    kesim = min(bugun, ay_son)
    gy_kesim = kesim.replace(year=yil - 1) if not (kesim.month == 2 and kesim.day == 29) else date(yil - 1, 2, 28)
    toplam_gun = is_gunu_say(ay_bas, ay_son, calisma_gunu)
    gecen_gun = is_gunu_say(ay_bas, kesim, calisma_gunu) if bugun >= ay_bas else 0
    kalan_gun = toplam_gun - gecen_gun
    ay_ici = bugun < ay_son

    def anahtar(x, kirilim):
        return tuple(x[k] for k in kirilim)

    sonuc = OrderedDict()
    kirilimlar = [("Genel", [])] + [(BOYUT_AD[k], [k]) for k in hb] + ([("Temsilci × Ürün Grubu", ["temsilci", "grup"])]
                                                                      if {"temsilci", "grup"} <= set(boyutlar) else [])
    for ad, kir in kirilimlar:
        tablo = defaultdict(lambda: {"ay": 0.0, "ytd": 0.0, "gy_ay": 0.0, "gy_ytd": 0.0, "hedef_ay": None, "hedef_ytd": None})
        for x in satislar:
            k = anahtar(x, kir)
            t = x["tarih"]
            if ay_bas <= t <= kesim:
                tablo[k]["ay"] += x["tutar"]
            if date(yil, 1, 1) <= t <= kesim:
                tablo[k]["ytd"] += x["tutar"]
            if date(yil - 1, aynum, 1) <= t <= gy_kesim:
                tablo[k]["gy_ay"] += x["tutar"]
            if date(yil - 1, 1, 1) <= t <= gy_kesim:
                tablo[k]["gy_ytd"] += x["tutar"]
        # Hedef: kırılım boyutları hedefin verildiği boyutların alt kümesiyse toplanabilir
        if set(kir) <= set(hb):
            idx = [hb.index(k) for k in kir]
            for (h_ay, *deger), v in hedef.items():
                k = tuple(deger[i] for i in idx)
                if h_ay == ay:
                    tablo[k]["hedef_ay"] = (tablo[k]["hedef_ay"] or 0) + v
                if h_ay[:4] == ay[:4] and h_ay <= ay:
                    tablo[k]["hedef_ytd"] = (tablo[k]["hedef_ytd"] or 0) + v
        for k, v in tablo.items():
            h = v["hedef_ay"]
            v["oran"] = v["ay"] / h if h else None
            v["oran_ytd"] = v["ytd"] / v["hedef_ytd"] if v["hedef_ytd"] else None
            v["gy_degisim"] = (v["ay"] / v["gy_ay"] - 1) if v["gy_ay"] else None
            v["gy_degisim_ytd"] = (v["ytd"] / v["gy_ytd"] - 1) if v["gy_ytd"] else None
            if ay_ici and gecen_gun:
                v["beklenen"] = h * gecen_gun / toplam_gun if h else None
                v["tahmin"] = v["ay"] / gecen_gun * toplam_gun
                v["gereken_gunluk"] = max(0.0, (h - v["ay"]) / kalan_gun) if h and kalan_gun else None
            else:
                v["beklenen"], v["tahmin"], v["gereken_gunluk"] = h, v["ay"], None
            v["tahmin_oran"] = v["tahmin"] / h if h else None
            v["prim_carpani"] = prim_carpani(v["oran"], kademeler) if kademeler and v["oran"] is not None and not ay_ici else None
        sonuc[ad] = (kir, OrderedDict(sorted(tablo.items(), key=lambda i: -(i[1]["ay"]))))
    return {"tablolar": sonuc, "toplam_gun": toplam_gun, "gecen_gun": gecen_gun, "kalan_gun": kalan_gun, "ay_ici": ay_ici,
            "kesim": kesim, "gy_kesim": gy_kesim}


def calistir(satis_yolu: Path, hedef_yolu: Path, cikti: Path, ay: str | None = None, bugun: date | None = None,
             calisma_gunu: int = 5, kademeler: list[tuple[float, float]] | None = None) -> dict:
    satislar, boyutlar = satislar_oku(satis_yolu)
    hedef, hb = hedefler_oku(hedef_yolu, boyutlar)
    if not satislar:
        raise SystemExit("Satış satırı yok.")
    ay = ay or max(f"{x['tarih'].year}-{x['tarih'].month:02d}" for x in satislar)
    bugun = bugun or date.today()
    h = hesapla(satislar, boyutlar, hedef, hb, ay, bugun, calisma_gunu, kademeler or [])
    h["ay"], h["boyutlar"], h["hb"] = ay, boyutlar, hb
    h["trend"] = trend(satislar, hedef, hb, ay)
    _rapor(h, cikti, kademeler)
    return h


def trend(satislar, hedef, hb, ay) -> list[tuple[str, float, float | None, float]]:
    """Yılın ayları: (ay, gerçekleşen, hedef, GY aynı ay)."""
    yil = int(ay[:4])
    sonuc = []
    for m in range(1, int(ay[5:]) + 1):
        a = f"{yil}-{m:02d}"
        g = sum(x["tutar"] for x in satislar if (x["tarih"].year, x["tarih"].month) == (yil, m))
        gy = sum(x["tutar"] for x in satislar if (x["tarih"].year, x["tarih"].month) == (yil - 1, m))
        hd = sum(v for k, v in hedef.items() if k[0] == a) or None
        sonuc.append((a, g, hd, gy))
    return sonuc


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
PARA = "#,##0"


def renk(oran: float | None) -> PatternFill | None:
    if oran is None:
        return None
    return PatternFill("solid", fgColor="E3F5E1" if oran >= 1 else "FFF4CE" if oran >= 0.9 else "FDE2E1")


def _rapor(h, cikti, kademeler):
    wb = Workbook()
    ilk = True
    for ad, (kir, tablo) in h["tablolar"].items():
        ws = wb.active if ilk else wb.create_sheet(ad[:30])
        if ilk:
            ws.title = "Genel"
        bas = [BOYUT_AD[k] for k in kir] or ["Kapsam"]
        ws.append([f"Satış hedef-gerçekleşme · {h['ay']} · {h['kesim']:%d.%m.%Y} itibarıyla"
                   + (f" · {h['gecen_gun']}/{h['toplam_gun']} iş günü geçti" if h["ay_ici"] else " · ay kapandı")])
        ws["A1"].font = Font(bold=True, size=12)
        sutunlar = bas + ["Ay Gerçekleşen", "Ay Hedefi", "Gerçekleşme %", "GY Aynı Dönem", "GY'ye Göre %"]
        if h["ay_ici"]:
            sutunlar += ["Bugüne Beklenen", "Ay Sonu Tahmini", "Tahmin / Hedef", "Gereken Günlük"]
        sutunlar += ["YTD Gerçekleşen", "YTD Hedef", "YTD %", "GY YTD", "YTD GY'ye Göre %"]
        if kademeler and not h["ay_ici"]:
            sutunlar += ["Prim Çarpanı"]
        ws.append(sutunlar)
        for c in ws[2]:
            c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI
        for k, v in tablo.items():
            satir = list(k) or ["Toplam"]
            satir += [v["ay"], v["hedef_ay"], v["oran"], v["gy_ay"] or None, v["gy_degisim"]]
            if h["ay_ici"]:
                satir += [v["beklenen"], v["tahmin"], v["tahmin_oran"], v["gereken_gunluk"]]
            satir += [v["ytd"], v["hedef_ytd"], v["oran_ytd"], v["gy_ytd"] or None, v["gy_degisim_ytd"]]
            if kademeler and not h["ay_ici"]:
                satir += [v["prim_carpani"]]
            ws.append(satir)
            n = ws.max_row
            for j, ad_s in enumerate(sutunlar, 1):
                if "%" in ad_s or "Tahmin / Hedef" in ad_s:
                    ws.cell(n, j).number_format = "0.0%"
                elif j > len(bas) and ad_s != "Prim Çarpanı":
                    ws.cell(n, j).number_format = PARA
            ws.cell(n, len(bas) + 3).fill = renk(v["oran"]) or PatternFill()
            if h["ay_ici"]:
                ws.cell(n, len(bas) + 8).fill = renk(v["tahmin_oran"]) or PatternFill()
        for j in range(1, len(sutunlar) + 1):
            ws.column_dimensions[get_column_letter(j)].width = 16 if j <= len(bas) else 14
        ws.freeze_panes = ws.cell(3, len(bas) + 1)
        if ilk:
            ws.append([])
            ws.append(["Ay", "Gerçekleşen", "Hedef", "GY Aynı Ay"])
            bas_satir = ws.max_row
            for c in ws[bas_satir]:
                c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI
            for a, g, hd, gy in h["trend"]:
                ws.append([a, g, hd, gy or None])
                for j in (2, 3, 4):
                    ws.cell(ws.max_row, j).number_format = PARA
            g = BarChart()
            g.title, g.height, g.width = "Aylık gerçekleşen, hedef ve geçen yıl", 8, 22
            g.add_data(Reference(ws, min_col=2, max_col=4, min_row=bas_satir, max_row=ws.max_row), titles_from_data=True)
            g.set_categories(Reference(ws, min_col=1, min_row=bas_satir + 1, max_row=ws.max_row))
            ws.add_chart(g, f"A{ws.max_row + 2}")
        ilk = False
    b = wb.create_sheet("Bilgi")
    for s in [["Dönem", f"{h['ay']} · kesim tarihi {h['kesim']:%d.%m.%Y} · GY karşılaştırması {h['gy_kesim']:%d.%m.%Y}'e kadar (aynı gün sayısı)"],
              ["İş günü", f"ayda {h['toplam_gun']} iş günü; {h['gecen_gun']} geçti, {h['kalan_gun']} kaldı (resmî tatiller hariç tutulmadı)"],
              ["Ay sonu tahmini", "bugüne kadarki gerçekleşen ÷ geçen iş günü × ayın iş günü (doğrusal run-rate)"],
              ["Gereken günlük", "(ay hedefi − gerçekleşen) ÷ kalan iş günü"],
              ["Hedef kırılımı", "hedefler yalnız hedef dosyasındaki boyutlarda (" + ", ".join(BOYUT_AD[k] for k in h["hb"]) +
                                 ") gösterilir; diğer kırılımlarda yalnız gerçekleşen ve GY karşılaştırması vardır"],
              ["Renk", "yeşil ≥ %100 · sarı %90–99 · kırmızı < %90"],
              ["İadeler", "Tür sütununda 'İade' yazan satırlar düşülür; tutarlar KDV hariç net satış olmalıdır"]]:
        b.append(s)
    if kademeler:
        b.append(["Prim kademeleri", " · ".join(f"≥ %{e:g} → {k:g} kat" for e, k in sorted(kademeler)) + " (yalnız kapanmış ay için)"])
    b.column_dimensions["A"].width = 18
    b.column_dimensions["B"].width = 130
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ornek = BURASI / "ornek_veri"
    ap = argparse.ArgumentParser(description="Temsilci, bölge ve ürün grubu bazında satış hedef-gerçekleşme raporu üretir.")
    ap.add_argument("--satislar", type=Path, default=ornek / "satislar.csv",
                    help="Satışlar: Tarih, Net Tutar [, Temsilci, Bölge, Ürün Grubu, Tür]")
    ap.add_argument("--hedefler", type=Path, default=ornek / "hedefler.csv",
                    help="Hedefler: Ay, [Temsilci, Bölge, Ürün Grubu,] Hedef — veya Temsilci, 2026-01, 2026-02 …")
    ap.add_argument("--ay", help="Rapor ayı YYYY-AA (varsayılan satışlardaki son ay)")
    ap.add_argument("--bugun", help="Kesim tarihi GG.AA.YYYY (ay içi rapor için; varsayılan bugün)")
    ap.add_argument("--calisma-gunu", type=int, default=5, choices=[5, 6], help="Haftalık iş günü (varsayılan 5)")
    ap.add_argument("--prim", help="Prim kademeleri, ör. 90=0.5,100=1,110=1.5 (gerçekleşme %% → çarpan)")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "hedef_gerceklesme.xlsx")
    a = ap.parse_args(argv)
    bugun = tarih(a.bugun) if a.bugun else None
    if a.satislar == ornek / "satislar.csv":
        a.ay = a.ay or "2026-10"
        bugun = bugun or date(2026, 10, 15)
    kademeler = []
    for p in (a.prim or "").split(","):
        if "=" in p:
            e, _, k = p.partition("=")
            kademeler.append((sayi(e), sayi(k)))
    s = calistir(a.satislar, a.hedefler, a.cikti, a.ay, bugun, a.calisma_gunu, kademeler)
    g = next(iter(s["tablolar"]["Genel"][1].values()))
    print(f"[OK] {s['ay']}: gerçekleşen {tl(g['ay'])} / hedef {tl(g['hedef_ay'] or 0)}"
          + (f" (%{tl(g['oran'] * 100, 1)})" if g["oran"] is not None else "")
          + (f" · ay sonu tahmini {tl(g['tahmin'])} · gereken günlük {tl(g['gereken_gunluk'] or 0)}" if s["ay_ici"] else "")
          + (f" · GY'ye göre %{tl(g['gy_degisim'] * 100, 1)}" if g["gy_degisim"] is not None else ""))
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
