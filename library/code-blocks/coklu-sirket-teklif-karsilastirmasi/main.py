"""
Çoklu Şirket Teklif Karşılaştırması — Workers / Workless kod bloğu
Sigorta › Teknik ve Operasyon › Sigorta Uzmanı

Farklı sigorta şirketlerinden alınan teklifleri teminat, limit, muafiyet ve prim bazında müşteriye sunulacak tek
tabloya döker:
  - Teminat adları şirketlere göre farklı yazılsa da ("Yangın (bina)" / "Yangın bina") müşterinin talep listesiyle
    eşleştirilir; talepte olmayan teminatlar "ek teminat" olarak gösterilir.
  - Her teminat için limit istenen limitin altında mı, muafiyet azami muafiyetin üstünde mi kontrol edilir. Farklı
    yazılmış muafiyetler ("%2, en az 10.000 TL", "5.000 TL") istenen limit tutarındaki bir hasar üzerinden TL'ye
    çevrilerek karşılaştırılır; "7 gün" gibi süre muafiyetleri karşılaştırılmaz, gösterilir.
  - Zorunlu teminatları eksiksiz ve uygun olan teklifler "Uygun"; diğerleri nedenleriyle listelenir.
  - Geçerlilik tarihi geçmiş teklifler, özel şartlar ve teminat notları (koasürans gibi) uyarı olarak çıkar.
  - Uygun teklifler brüt prime göre sıralanır; en düşük primli ve tüm talebi karşılayan en düşük primli teklif ayrıca
    belirtilir. Seçim müşterinindir.
İnternete bağlanmaz.

Kullanım:
    python main.py                                                     # örnek işyeri paket sigortası, 4 kurgusal şirket
    python main.py --teklifler teklifler.xlsx --ozet teklif_ozeti.xlsx --talep talep.xlsx --tarih 09.10.2026
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
SIFIR = Decimal(0)
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
DOLGU_KELIMELER = {"ve", "ile", "teminati", "teminat", "sigortasi", "dahil", "the"}

TEKLIF_SUTUNLARI = {"sirket": ("sirket", "sigorta sirketi", "firma"), "teminat": ("teminat", "teminat adi"), "limit": ("limit", "teminat limiti", "bedel"),
                    "muafiyet": ("muafiyet", "tenzili muafiyet"), "not": ("not", "aciklama", "ozel sart")}
OZET_SUTUNLARI = {"sirket": TEKLIF_SUTUNLARI["sirket"], "net": ("net prim",), "brut": ("brut prim", "toplam prim", "prim"),
                  "odeme": ("odeme plani", "odeme", "taksit"), "gecerlilik": ("gecerlilik tarihi", "gecerlilik", "son gecerlilik"),
                  "ozel": ("ozel sart", "ozel sartlar", "not")}
TALEP_SUTUNLARI = {"teminat": ("teminat",), "limit": ("istenen limit", "limit"), "muafiyet": ("azami muafiyet", "muafiyet"),
                   "zorunlu": ("zorunlu", "zorunlu mu")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(str(x or "").strip(), f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif s.count(".") > 1 or (s.count(".") == 1 and len(s.split(".")[1]) == 3):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tl(x) -> str:
    return "—" if x is None else f"{x:,.0f}".replace(",", ".")


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} for r in s[bi + 1:]]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


# ----------------------------------------------------------------------------
# Muafiyet
# ----------------------------------------------------------------------------

@dataclass
class Muafiyet:
    metin: str
    oran: Decimal | None = None
    en_az: Decimal | None = None
    en_cok: Decimal | None = None
    sabit: Decimal | None = None
    sure: str | None = None

    @property
    def yok(self) -> bool:
        return not self.metin.strip() or katla(self.metin) in ("yok", "0", "muafiyetsiz")

    def tutar(self, hasar: Decimal) -> Decimal | None:
        """Verilen hasar tutarında muafiyet (TL); süre muafiyetinde None."""
        if self.yok:
            return SIFIR
        if self.sure:
            return None
        if self.sabit is not None:
            return self.sabit
        t = hasar * self.oran if self.oran is not None else SIFIR
        if self.en_az is not None:
            t = max(t, self.en_az)
        if self.en_cok is not None:
            t = min(t, self.en_cok)
        return t


def muafiyet_coz(metin) -> Muafiyet:
    ham = str(metin or "").strip()
    k = ham.lower().translate(_TR)
    bul = lambda desen: (lambda m: para(m.group(1).rstrip(".,")) if m else None)(re.search(desen, k))  # noqa: E731
    m = Muafiyet(ham)
    sure = re.search(r"(\d+)\s*(gun|saat|hafta)", k)
    if sure:
        m.sure = sure.group(0)
        return m
    oran = bul(r"%\s*([\d.,]+)")
    m.oran = oran / 100 if oran is not None else None
    m.en_az, m.en_cok = bul(r"(?:en az|asgari|min)\w*\s*([\d.,]+)"), bul(r"(?:en fazla|azami|maks)\w*\s*([\d.,]+)")
    if oran is None:
        m.sabit = bul(r"([\d][\d.,]*)")
    return m


# ----------------------------------------------------------------------------
# Eşleştirme ve değerlendirme
# ----------------------------------------------------------------------------

def kelimeler(ad: str) -> set[str]:
    return {w for w in katla(ad).split() if w not in DOLGU_KELIMELER}


def benzerlik(a: str, b: str) -> float:
    x, y = kelimeler(a), kelimeler(b)
    if not x or not y:
        return 0.0
    if x == y:
        return 1.0
    return len(x & y) / len(x | y)


@dataclass
class Talep:
    teminat: str
    limit: Decimal | None
    muafiyet: Muafiyet
    zorunlu: bool


@dataclass
class Hucre:
    teminat: str
    limit: Decimal | None
    muafiyet: Muafiyet
    not_: str
    durum: str = "Uygun"
    sorunlar: list = field(default_factory=list)


@dataclass
class Teklif:
    sirket: str
    net: Decimal | None = None
    brut: Decimal | None = None
    odeme: str = ""
    gecerlilik: date | None = None
    ozel: str = ""
    hucreler: dict = field(default_factory=dict)      # talep teminat adı → Hucre
    ekler: list = field(default_factory=list)          # talepte olmayan teminatlar
    uyarilar: list = field(default_factory=list)
    eksik_zorunlu: list = field(default_factory=list)
    eksik_istege: list = field(default_factory=list)
    uygunsuz: list = field(default_factory=list)
    uygun: bool = True


def esle(ad: str, talepler: list[Talep]) -> Talep | None:
    en = max(talepler, key=lambda t: benzerlik(ad, t.teminat), default=None)
    return en if en and benzerlik(ad, en.teminat) >= 0.6 else None


def degerlendir(teklif_satirlari: list[dict], ozet_satirlari: list[dict], talepler: list[Talep], bugun: date) -> list[Teklif]:
    teklifler: dict[str, Teklif] = {}
    for r in teklif_satirlari:
        s = str(r.get("sirket") or "").strip()
        if not s or not r.get("teminat"):
            continue
        t = teklifler.setdefault(s, Teklif(s))
        ad = str(r["teminat"]).strip()
        h = Hucre(ad, para(r.get("limit")), muafiyet_coz(r.get("muafiyet")), str(r.get("not") or "").strip())
        talep = esle(ad, talepler) if talepler else None
        if talep is None:
            if talepler:
                t.ekler.append(h)
                continue
            t.hucreler[ad] = h
            continue
        if talep.teminat in t.hucreler:
            t.uyarilar.append(f"'{ad}' ve '{t.hucreler[talep.teminat].teminat}' aynı talebe ({talep.teminat}) eşleşti; ilki kullanıldı")
            continue
        t.hucreler[talep.teminat] = h
    for r in ozet_satirlari:
        s = str(r.get("sirket") or "").strip()
        if not s:
            continue
        t = teklifler.setdefault(s, Teklif(s))
        t.net, t.brut, t.odeme = para(r.get("net")), para(r.get("brut")), str(r.get("odeme") or "").strip()
        t.gecerlilik, t.ozel = tarih(r.get("gecerlilik")), str(r.get("ozel") or "").strip()

    for t in teklifler.values():
        for talep in talepler:
            h = t.hucreler.get(talep.teminat)
            if h is None:
                (t.eksik_zorunlu if talep.zorunlu else t.eksik_istege).append(talep.teminat)
                continue
            if talep.limit is not None and (h.limit is None or h.limit < talep.limit):
                h.sorunlar.append(f"limit {tl(h.limit)} < istenen {tl(talep.limit)}")
            if not talep.muafiyet.yok:
                ref = talep.limit or h.limit or Decimal(100000)
                istenen, verilen = talep.muafiyet.tutar(ref), h.muafiyet.tutar(min(ref, h.limit or ref))
                if verilen is None or istenen is None:
                    h.sorunlar.append(f"muafiyet '{h.muafiyet.metin}' azami muafiyetle ('{talep.muafiyet.metin}') karşılaştırılamadı")
                    h.durum = "Kontrol"
                elif verilen > istenen:
                    h.sorunlar.append(f"muafiyet {h.muafiyet.metin} ({tl(verilen)} TL) > azami {talep.muafiyet.metin} ({tl(istenen)} TL; {tl(ref)} TL hasarda)")
            if any(not x.startswith("muafiyet '") for x in h.sorunlar):
                h.durum = "Uygun değil"
                (t.uygunsuz).append((talep.teminat, talep.zorunlu, "; ".join(h.sorunlar)))
            if h.not_:
                t.uyarilar.append(f"{talep.teminat}: {h.not_}")
        for h in t.ekler:
            t.uyarilar.append(f"Ek teminat (talepte yok): {h.teminat} {tl(h.limit)} TL" + (f", muafiyet {h.muafiyet.metin}" if h.muafiyet.metin else ""))
        if t.brut is None:
            t.uyarilar.append("Brüt prim yok (teklif özetinde şirket bulunamadı)")
        if t.gecerlilik and t.gecerlilik < bugun:
            t.uyarilar.append(f"Teklifin geçerlilik süresi {t.gecerlilik:%d.%m.%Y} tarihinde dolmuş; şirketten yenisi istenmeli")
        if t.ozel:
            t.uyarilar.append(f"Özel şart: {t.ozel}")
        t.uygun = not t.eksik_zorunlu and not any(z for _, z, _ in t.uygunsuz) and not (t.gecerlilik and t.gecerlilik < bugun) and t.brut is not None
    return sorted(teklifler.values(), key=lambda t: (not t.uygun, t.brut if t.brut is not None else Decimal("1e18")))


def oneriler(teklifler: list[Teklif], talepler: list[Talep]) -> dict:
    uygun = [t for t in teklifler if t.uygun]
    tam = [t for t in uygun if not t.eksik_istege and not t.uygunsuz]
    return {"en_dusuk": uygun[0] if uygun else None, "tam_en_dusuk": tam[0] if tam else None, "uygun_sayisi": len(uygun)}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
YESIL, SARI, KIRMIZI, GRI = (PatternFill("solid", fgColor=c) for c in ("E3F4E1", "FFF4CE", "FDE2E1", "EEEEEE"))
UST = Alignment(vertical="top", wrap_text=True)


def rapor_yaz(cikti: Path, teklifler: list[Teklif], talepler: list[Talep], o: dict, bugun: date) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Karşılaştırma"
    sirketler = [t.sirket for t in teklifler]
    ws.append(["Teminat", "İstenen Limit", "Azami Muafiyet", *[x for s in sirketler for x in (f"{s}\nLimit", f"{s}\nMuafiyet")]])
    for h in ws[1]:
        h.fill, h.font, h.alignment = BASLIK_DOLGU, BASLIK_YAZI, Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 34, 14, 14
    for j in range(4, 4 + 2 * len(sirketler)):
        ws.column_dimensions[get_column_letter(j)].width = 15
    ws.freeze_panes = "D2"
    satirlar = [(tp.teminat + ("" if tp.zorunlu else " (isteğe bağlı)"), tp.limit, tp.muafiyet.metin, tp.teminat) for tp in talepler]
    if not talepler:
        adlar = list(dict.fromkeys(a for t in teklifler for a in t.hucreler))
        satirlar = [(a, None, "", a) for a in adlar]
    for ad, lim, muaf, anahtar in satirlar:
        ws.append([ad, lim and float(lim), muaf or "—"])
        ws.cell(ws.max_row, 2).number_format = "#,##0"
        for i, t in enumerate(teklifler):
            h = t.hucreler.get(anahtar)
            c1, c2 = ws.cell(ws.max_row, 4 + 2 * i), ws.cell(ws.max_row, 5 + 2 * i)
            if h is None:
                c1.value, c2.value = "YOK", ""
                c1.fill = c2.fill = GRI if "isteğe" in ad else KIRMIZI
                continue
            c1.value, c2.value = h.limit and float(h.limit), h.muafiyet.metin or "Yok"
            c1.number_format = "#,##0"
            kotu = SARI if "isteğe" in ad else KIRMIZI
            c1.fill = kotu if any(x.startswith("limit") for x in h.sorunlar) else YESIL
            muaf_kotu = any(x.startswith("muafiyet ") and "karşılaştırılamadı" not in x for x in h.sorunlar)
            c2.fill = kotu if muaf_kotu else SARI if h.durum == "Kontrol" else YESIL
    for etiket, deger, fmt in (("Ek teminatlar", lambda t: "; ".join(f"{h.teminat} {tl(h.limit)}" for h in t.ekler) or "—", None),
                               ("Net prim", lambda t: t.net and float(t.net), "#,##0"), ("Brüt prim", lambda t: t.brut and float(t.brut), "#,##0"),
                               ("Ödeme planı", lambda t: t.odeme or "—", None),
                               ("Geçerlilik", lambda t: f"{t.gecerlilik:%d.%m.%Y}" if t.gecerlilik else "—", None),
                               ("Sonuç", lambda t: "Uygun" if t.uygun else "Uygun değil", None)):
        ws.append([etiket, None, None])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        for i, t in enumerate(teklifler):
            c = ws.cell(ws.max_row, 4 + 2 * i, deger(t))
            ws.merge_cells(start_row=ws.max_row, start_column=4 + 2 * i, end_row=ws.max_row, end_column=5 + 2 * i)
            c.alignment = UST
            if fmt:
                c.number_format = fmt
            if etiket == "Sonuç":
                c.fill, c.font = (YESIL if t.uygun else KIRMIZI), Font(bold=True)
            if etiket == "Geçerlilik" and t.gecerlilik and t.gecerlilik < bugun:
                c.fill = KIRMIZI

    dg = wb.create_sheet("Değerlendirme")
    dg.append(["Sıra", "Şirket", "Brüt Prim", "Sonuç", "Eksik Zorunlu Teminat", "Uygun Olmayan Teminatlar", "Eksik İsteğe Bağlı", "Uyarılar", "Müşteri Tercihi"])
    for h in dg[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for j, w in enumerate((5, 18, 12, 11, 26, 50, 26, 60, 16), 1):
        dg.column_dimensions[get_column_letter(j)].width = w
    for n, t in enumerate(teklifler, 1):
        dg.append([n if t.uygun else "—", t.sirket, t.brut and float(t.brut), "Uygun" if t.uygun else "Uygun değil", ", ".join(t.eksik_zorunlu) or "—",
                   "\n".join(f"{a}{'' if z else ' (isteğe bağlı)'}: {s}" for a, z, s in t.uygunsuz) or "—", ", ".join(t.eksik_istege) or "—",
                   "\n".join(t.uyarilar) or "—", ""])
        dg.cell(dg.max_row, 3).number_format = "#,##0"
        dg.cell(dg.max_row, 4).fill = YESIL if t.uygun else KIRMIZI
        dg.cell(dg.max_row, 9).fill = SARI
        for c in dg[dg.max_row]:
            c.alignment = UST
    dg.append([])
    for et, t in (("En düşük primli uygun teklif", o["en_dusuk"]), ("Tüm talebi karşılayan en düşük primli teklif", o["tam_en_dusuk"])):
        dg.append(["", et, t.sirket if t else "yok", f"{tl(t.brut)} TL" if t else ""])
        dg.cell(dg.max_row, 2).font = Font(bold=True)
    dg.append(["", "Not", "Karşılaştırma bilgi amaçlıdır; seçim müşteriye aittir. Poliçe genel ve özel şartları, klozlar ve şirketlerin yazılı teklifleri esastır."])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

    md = cikti.with_suffix(".md")
    sat = [f"# Sigorta Teklifleri Karşılaştırması — {bugun:%d.%m.%Y}", "", "| Teminat | İstenen | " + " | ".join(sirketler) + " |",
           "|---|---:|" + "---|" * len(sirketler)]
    for ad, lim, muaf, anahtar in satirlar:
        hucre = []
        for t in teklifler:
            h = t.hucreler.get(anahtar)
            hucre.append("**YOK**" if h is None else f"{tl(h.limit)}" + (f" / {h.muafiyet.metin}" if h.muafiyet.metin else "") + (" ⚠" if h.sorunlar else ""))
        sat.append(f"| {ad} | {tl(lim)}{(' / ' + muaf) if muaf else ''} | " + " | ".join(hucre) + " |")
    sat.append("| **Brüt prim (TL)** | | " + " | ".join(f"**{tl(t.brut)}**" for t in teklifler) + " |")
    sat.append("| Ödeme planı | | " + " | ".join(t.odeme or "—" for t in teklifler) + " |")
    sat.append("| **Sonuç** | | " + " | ".join("Uygun" if t.uygun else "Uygun değil" for t in teklifler) + " |")
    sat += ["", "Hücreler: limit / muafiyet (TL). ⚠ = talebinizi karşılamıyor.", ""]
    for t in teklifler:
        notlar = [f"Eksik zorunlu teminat: {', '.join(t.eksik_zorunlu)}"] if t.eksik_zorunlu else []
        notlar += [f"{a}: {s}" for a, _, s in t.uygunsuz] + t.uyarilar
        if notlar:
            sat += [f"**{t.sirket}**", ""] + [f"- {x}" for x in notlar] + [""]
    if o["en_dusuk"]:
        sat.append(f"En düşük primli uygun teklif: **{o['en_dusuk'].sirket}** ({tl(o['en_dusuk'].brut)} TL).")
    if o["tam_en_dusuk"]:
        sat.append(f"Tüm talebi karşılayan en düşük primli teklif: **{o['tam_en_dusuk'].sirket}** ({tl(o['tam_en_dusuk'].brut)} TL).")
    sat += ["", "> Bilgi amaçlıdır; seçim size aittir. Poliçe genel ve özel şartları ile şirketlerin yazılı teklifleri esastır."]
    md.write_text("\n".join(sat) + "\n", encoding="utf-8")
    return md


def calistir(teklif_yolu: Path, ozet_yolu: Path | None, talep_yolu: Path | None, cikti: Path, bugun: date | None = None) -> dict:
    bugun = bugun or date.today()
    satirlar = kayitlar(teklif_yolu, TEKLIF_SUTUNLARI, ("sirket", "teminat"))
    ozet = kayitlar(ozet_yolu, OZET_SUTUNLARI, ("sirket", "brut")) if ozet_yolu else []
    talepler = [Talep(str(r["teminat"]).strip(), para(r.get("limit")), muafiyet_coz(r.get("muafiyet")),
                      katla(r.get("zorunlu")) not in ("hayir", "h", "no", "istege bagli", "opsiyonel"))
                for r in (kayitlar(talep_yolu, TALEP_SUTUNLARI, ("teminat",)) if talep_yolu else []) if r.get("teminat")]
    teklifler = degerlendir(satirlar, ozet, talepler, bugun)
    o = oneriler(teklifler, talepler)
    md = rapor_yaz(cikti, teklifler, talepler, o, bugun)
    return {"teklifler": teklifler, "talepler": talepler, "oneri": o, "md": md}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Sigorta şirketlerinden alınan teklifleri teminat, limit, muafiyet ve prim bazında karşılaştırır.")
    p.add_argument("--teklifler", type=Path, default=ORNEK / "teklifler.csv", help="Şirket, Teminat, Limit, Muafiyet, Not (.xlsx/.csv)")
    p.add_argument("--ozet", type=Path, help="Şirket, Net Prim, Brüt Prim, Ödeme Planı, Geçerlilik Tarihi, Özel Şart")
    p.add_argument("--talep", type=Path, help="Müşterinin talebi: Teminat, İstenen Limit, Azami Muafiyet, Zorunlu")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY (varsayılan: bugün; örnek veride 08.10.2026)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "teklif_karsilastirma.xlsx")
    a = p.parse_args(argv)
    ornek = a.teklifler == ORNEK / "teklifler.csv"
    ozet = a.ozet or (ORNEK / "teklif_ozeti.csv" if ornek else None)
    talep = a.talep or (ORNEK / "talep.csv" if ornek else None)
    bugun = tarih(a.tarih) if a.tarih else (date(2026, 10, 8) if ornek else None)
    if a.tarih and bugun is None:
        print(f"[X] Tarih GG.AA.YYYY olmalı: {a.tarih}")
        return 2
    for y in (a.teklifler, ozet, talep):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        s = calistir(a.teklifler, ozet, talep, a.cikti, bugun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    for t in s["teklifler"]:
        print(f"[{'OK' if t.uygun else 'X'}] {t.sirket}: brüt {tl(t.brut)} TL · " + ("uygun" if t.uygun else "uygun değil: " +
              "; ".join([f"eksik {', '.join(t.eksik_zorunlu)}"] * bool(t.eksik_zorunlu) + [a for a, z, _ in t.uygunsuz if z] +
                        ["süresi dolmuş"] * bool(t.gecerlilik and t.gecerlilik < (bugun or date.today())))))
    o = s["oneri"]
    if o["en_dusuk"]:
        print(f"[i] En düşük primli uygun teklif: {o['en_dusuk'].sirket}" + (f" · tüm talebi karşılayan: {o['tam_en_dusuk'].sirket}" if o["tam_en_dusuk"] else ""))
    print(f"[OK] Rapor: {a.cikti.resolve()} · {s['md'].resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
