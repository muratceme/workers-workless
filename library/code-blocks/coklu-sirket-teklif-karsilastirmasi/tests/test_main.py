import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "t.xlsx"
        cls.s = main.calistir(ORNEK / "teklifler.csv", ORNEK / "teklif_ozeti.csv", ORNEK / "talep.csv", cls.cikti, date(2026, 10, 8))
        cls.t = {t.sirket: t for t in cls.s["teklifler"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_siralama_ve_sonuc(self):
        self.assertEqual([(t.sirket, t.uygun) for t in self.s["teklifler"]],
                         [("Örnek A Sigorta", True), ("Örnek D Sigorta", True), ("Örnek C Sigorta", False), ("Örnek B Sigorta", False)])
        self.assertEqual(self.s["oneri"]["en_dusuk"].sirket, "Örnek A Sigorta")
        self.assertEqual(self.s["oneri"]["tam_en_dusuk"].sirket, "Örnek D Sigorta")

    def test_eslestirme(self):
        b = self.t["Örnek B Sigorta"]
        self.assertEqual(b.hucreler["Yangın (bina)"].teminat, "Yangın bina")
        self.assertEqual(b.hucreler["Yangın (emtia ve demirbaş)"].teminat, "Yangın emtia demirbaş")
        self.assertEqual([h.teminat for h in self.t["Örnek D Sigorta"].ekler], ["Elektronik cihaz"])
        self.assertEqual(self.t["Örnek A Sigorta"].eksik_istege, ["Kira kaybı", "İş durması"])

    def test_uygunsuzluklar(self):
        b = dict((a, s) for a, _, s in self.t["Örnek B Sigorta"].uygunsuz)
        self.assertIn("muafiyet %5 (550.000 TL) > azami %2 (220.000 TL", b["Deprem ve yanardağ püskürmesi"])
        self.assertEqual(b["Dahili su"], "limit 250.000 < istenen 500.000")
        c = self.t["Örnek C Sigorta"]
        self.assertIn("05.10.2026 tarihinde dolmuş", " ".join(c.uyarilar))
        self.assertIn("Koasürans %80", " ".join(self.t["Örnek A Sigorta"].uyarilar))

    def test_muafiyet(self):
        m = main.muafiyet_coz("%2, en az 10.000 TL")
        self.assertEqual((m.tutar(D(250000)), m.tutar(D(1000000))), (D(10000), D(20000)))
        self.assertEqual(main.muafiyet_coz("7 gün").tutar(D(1)), None)
        self.assertEqual(main.muafiyet_coz("").tutar(D(5)), 0)
        self.assertEqual(main.muafiyet_coz("%5 en fazla 1.000").tutar(D(100000)), D(1000))

    def test_cikti(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Karşılaştırma", "Değerlendirme"])
        ws = wb["Karşılaştırma"]
        self.assertEqual(ws["A1"].value, "Teminat")
        son = [r for r in ws.iter_rows(values_only=True) if r[0] == "Sonuç"][0]
        self.assertEqual((son[3], son[9]), ("Uygun", "Uygun değil"))
        md = self.s["md"].read_text(encoding="utf-8")
        self.assertIn("| Dahili su | 500.000 / 10.000 TL |", md)
        self.assertIn("250.000 / %2, en az 10.000 TL ⚠", md)

    def test_talepsiz(self):
        s = main.calistir(ORNEK / "teklifler.csv", ORNEK / "teklif_ozeti.csv", None, self.cikti, date(2026, 10, 1))
        self.assertTrue(all(t.uygun for t in s["teklifler"]))
        self.assertEqual(s["teklifler"][0].sirket, "Örnek C Sigorta")

    def test_cli(self):
        self.assertEqual(main.main(["--cikti", str(self.cikti)]), 0)
        self.assertEqual(main.main(["--tarih", "x", "--cikti", str(self.cikti)]), 2)


if __name__ == "__main__":
    unittest.main()
