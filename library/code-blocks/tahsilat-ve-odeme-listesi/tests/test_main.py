import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
B = date(2026, 10, 12)
g = lambda gun, ay=10: date(2026, ay, gun)   # noqa: E731


def k(cari, belge, vade, tutar):
    return {"cari": cari, "belge": belge, "vade": vade, "tutar": D(tutar), "aciklama": ""}


ALACAK = lambda: [k("X", "F1", g(5), 10000), k("Y", "F2", g(12), 20000), k("x", "F3", g(15), 5000), k("Z", "F4", g(25), 8000)]   # noqa: E731
BORC = lambda: [k("P", "B1", g(10), 15000), k("Q", "B2", g(13), 30000), k("P", "B3", g(18), 4000), k("R", "B4", g(1, 11), 9000)]   # noqa: E731


class ListeTesti(unittest.TestCase):
    def test_listeler(self):
        r = main.hazirla(ALACAK(), BORC(), B, kasa=D(10000))
        t, o = r["tahsilat"], r["odeme"]
        self.assertEqual([(x["belge"], x["durum"], x["gecikme"]) for x in t["satirlar"]], [("F1", "Vadesi geçmiş", 7), ("F2", "Bu dönem", 0), ("F3", "Bu dönem", 0)])
        self.assertEqual((t["gecmis"], t["donem"], t["toplam"], t["ileri"], o["gecmis"], o["donem"], o["toplam"], o["ileri"], r["son"], r["net"]), (D(10000), D(25000), D(35000), D(8000), D(15000), D(34000), D(49000), D(9000), g(18), D(-9000)))
        self.assertEqual([(c["cari"], c["kalem"], c["gecmis"], c["donem"], c["toplam"], c["en_eski"]) for c in t["cariler"]], [("Y", 1, D(0), D(20000), D(20000), g(12)), ("X", 2, D(10000), D(5000), D(15000), g(5))])
        self.assertEqual([(c["cari"], c["toplam"]) for c in o["cariler"]], [("Q", D(30000)), ("P", D(19000))])

    def test_gunluk_nakit(self):
        r = main.hazirla(ALACAK(), BORC(), B, kasa=D(10000))
        self.assertEqual((r["baslangic"], [(x["tarih"].day, x["gun"], x["tahsilat"], x["odeme"], x["net"], x["kumulatif"]) for x in r["gunluk"]]),
                         (D(-5000), [(12, "Pzt", D(20000), D(0), D(20000), D(15000)), (13, "Sal", D(0), D(30000), D(-30000), D(-15000)), (14, "Çar", D(0), D(0), D(0), D(-15000)), (15, "Per", D(5000), D(0), D(5000), D(-10000)),
                                     (16, "Cum", D(0), D(0), D(0), D(-10000)), (17, "Cmt", D(0), D(0), D(0), D(-10000)), (18, "Paz", D(0), D(4000), D(-4000), D(-14000))]))
        self.assertEqual((r["en_dusuk"]["tarih"], sorted((u["onem"], u["tur"]) for u in r["uyarilar"])), (g(13), [("Bilgi", "Hafta sonuna denk gelen vade"), ("Orta", "Vadesi geçmiş ödeme"), ("Yüksek", "Nakit açığı")]))
        r = main.hazirla(ALACAK(), BORC(), B, kasa=D(40000))
        self.assertEqual((r["baslangic"], r["gunluk"][1]["kumulatif"], [u["tur"] for u in r["uyarilar"] if u["onem"] == "Yüksek"]), (D(25000), D(15000), []))
        r = main.hazirla(ALACAK(), BORC(), B, kasa=D(20000))
        self.assertIn("13.10.2026 günü eksiye", [u for u in r["uyarilar"] if u["tur"] == "Nakit açığı"][0]["aciklama"])
        self.assertEqual([u["tur"] for u in main.hazirla(ALACAK(), BORC(), B)["uyarilar"] if u["onem"] == "Yüksek"], [])            # kasa verilmezse açık uyarısı üretilmez

    def test_donem_ve_sinirlar(self):
        r = main.hazirla(ALACAK(), BORC(), B, gun=1)
        self.assertEqual((r["son"], r["tahsilat"]["donem"], r["tahsilat"]["ileri"], r["odeme"]["donem"], len(r["gunluk"])), (B, D(20000), D(13000), D(0), 1))
        r = main.hazirla(ALACAK(), [], g(15, 11), gun=7)
        self.assertEqual(([x["gecikme"] for x in r["tahsilat"]["satirlar"]], [u["kim"] for u in r["uyarilar"] if u["tur"] == "30 günü aşan gecikmiş alacak"], r["odeme"]["toplam"]), ([41, 34, 31, 21], ["X · F1", "Y · F2", "x · F3"], D(0)))
        with self.assertRaises(ValueError):
            main.hazirla([], [], B)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "alacaklar.csv", ORNEK / "borclar.csv", cikti, B, 7, D(10000))
            self.assertEqual((r["tahsilat"]["toplam"], r["odeme"]["toplam"], r["net"], r["en_dusuk"]["kumulatif"], [(u["onem"], u["tur"]) for u in r["uyarilar"]]),
                             (D("35000.00"), D("49000.00"), D("-9000.00"), D("-15000.00"), [("Yüksek", "Nakit açığı"), ("Orta", "Vadesi geçmiş ödeme"), ("Bilgi", "Hafta sonuna denk gelen vade")]))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Tahsilat Listesi", "Ödeme Listesi", "Günlük Özet", "Cari Özeti", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            ay, by = Path(tmp) / "a.csv", Path(tmp) / "b.csv"
            ay.write_text("Müşteri;Fatura No;Vade Tarihi;Bakiye\nA;F1;12.10.2026;1.250,50\nB;F2;x;5\nC;F3;12.10.2026;0\n;F4;12.10.2026;5\n", encoding="utf-8")
            by.write_text("Tedarikçi;Vade;Tutar\nP;13.10.2026;500\n", encoding="utf-8")
            l, uy = main.oku(ay, "Alacak")
            self.assertEqual(([(x["cari"], x["belge"], x["vade"], x["tutar"]) for x in l], [u["tur"] for u in uy]), ([("A", "F1", B, D("1250.50"))], ["Satır okunamadı"] * 3))
            c = str(Path(tmp) / "x.xlsx")
            cik = ["--cikti", c]
            self.assertEqual(main.main(["--alacaklar", str(ay), "--borclar", str(by), "--tarih", "12.10.2026", "--gun", "14", "--kasa", "1.000,00"] + cik), 0)
            self.assertEqual(main.main(["--borclar", str(by)] + cik), 0)
            for ek in (["--tarih", "x"], ["--gun", "0"], ["--kasa", "x"]):
                self.assertEqual(main.main(["--alacaklar", str(ay)] + ek + cik), 2)
            self.assertEqual(main.main(["--alacaklar", str(Path(tmp) / "yok.csv")] + cik), 1)
            ay.write_text("Müşteri;Vade\nA;1\n", encoding="utf-8")
            self.assertEqual(main.main(["--alacaklar", str(ay)] + cik), 1)
            ay.write_text("Müşteri;Vade;Tutar\nA;x;1\n", encoding="utf-8")
            self.assertEqual(main.main(["--alacaklar", str(ay)] + cik), 1)
            self.assertEqual(main.main(cik), 0)


if __name__ == "__main__":
    unittest.main()
