"""
Tahsilat ve Ödeme Listesi — Workers / Workless kod bloğu
Muhasebe › Ön Muhasebe Elemanı

Açık alacak ve borç kalemlerinden haftalık tahsilat ve ödeme listesini, günlük nakit etkisini çıkarır.
  Dönem              = --tarih gününden başlayan --gun (7) gün
  Vadesi geçmiş      = vadesi dönem başından önce olan açık kalem;  gecikme = dönem başı − vade (gün)
  Bu dönem           = vadesi dönem içinde olan kalem
  İleri vadeli       = vadesi dönem sonundan sonra; listeye girmez, yalnız toplamı gösterilir
  Günlük net         = o gün vadesi gelen tahsilat − o gün vadesi gelen ödeme
  Kümülatif nakit    = başlangıç kasası (--kasa) − vadesi geçmiş ödemeler + günlük netlerin toplamı
                       (ihtiyatlı varsayım: gecikmiş borç hemen ödenir, gecikmiş alacak dönem içinde tahsil edilemez)
Liste vadelere dayanır; tahsilatın o gün gerçekleşeceğini göstermez. Tutarlar tek para biriminde olmalıdır; dövizli kalemleri kendi kurunuzla çevirip girin.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"cari": ("cari", "cari adi", "musteri", "tedarikci", "firma", "unvan"), "belge": ("belge no", "fatura no", "belge", "evrak no"), "vade": ("vade", "vade tarihi"), "tutar": ("tutar", "kalan tutar", "bakiye", "acik tutar"),
         "aciklama": ("aciklama", "not")}
Z, GUN = Decimal(0), timedelta(days=1)
GUNLER = ("Pzt", "Sal", "Çar", "Per", "Cum", "Cmt", "Paz")


def oku(yol: Path, ad: str) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("cari", "vade", "tutar")):
        c = metin(r["cari"])
        if not c and not metin(r["tutar"]):
            continue
        v, t = tarih(r["vade"]), para(r["tutar"])
        if not c or v is None or t is None or t <= 0:
            uy.append(uyari("Yüksek", "Satır okunamadı", f"{ad}: {c or '?'} (satır {r['_satir']})", f"Vade '{metin(r['vade'])}', tutar '{metin(r['tutar'])}'. Tutar pozitif olmalı. Kalem listeye alınmadı."))
            continue
        l.append({"cari": c, "belge": metin(r["belge"]), "vade": v, "tutar": t, "aciklama": metin(r["aciklama"])})
    return l, uy


def liste_kur(kalemler: list[dict], bas: date, son: date) -> dict:
    satirlar, ileri = [], Z
    for k in kalemler:
        if k["vade"] > son:
            ileri += k["tutar"]
            continue
        gecmis = k["vade"] < bas
        satirlar.append(k | {"durum": "Vadesi geçmiş" if gecmis else "Bu dönem", "gecikme": (bas - k["vade"]).days if gecmis else 0})
    satirlar.sort(key=lambda k: (k["vade"], k["cari"], k["belge"]))
    cari = {}
    for k in satirlar:
        o = cari.setdefault(katla(k["cari"]), {"cari": k["cari"], "kalem": 0, "gecmis": Z, "donem": Z, "en_eski": k["vade"]})
        o["kalem"] += 1
        o["gecmis" if k["durum"] == "Vadesi geçmiş" else "donem"] += k["tutar"]
        o["en_eski"] = min(o["en_eski"], k["vade"])
    gecmis, donem = sum((k["tutar"] for k in satirlar if k["durum"] == "Vadesi geçmiş"), Z), sum((k["tutar"] for k in satirlar if k["durum"] == "Bu dönem"), Z)
    return {"satirlar": satirlar, "cariler": sorted(({**o, "toplam": o["gecmis"] + o["donem"]} for o in cari.values()), key=lambda o: (-(o["gecmis"] + o["donem"]), o["cari"])), "gecmis": gecmis, "donem": donem, "ileri": ileri,
            "toplam": gecmis + donem}


def hazirla(alacaklar: list[dict], borclar: list[dict], bas: date, gun: int = 7, kasa: Decimal | None = None) -> dict:
    if not alacaklar and not borclar:
        raise ValueError("Listelenecek alacak veya borç kalemi yok.")
    son, uy = bas + (gun - 1) * GUN, []
    t, o = liste_kur(alacaklar, bas, son), liste_kur(borclar, bas, son)
    baslangic = (kasa or Z) - o["gecmis"]
    gunluk, kum, en_dusuk = [], baslangic, None
    for i in range(gun):
        g = bas + i * GUN
        th, od = sum((k["tutar"] for k in t["satirlar"] if k["vade"] == g), Z), sum((k["tutar"] for k in o["satirlar"] if k["vade"] == g), Z)
        kum += th - od
        gunluk.append({"tarih": g, "gun": GUNLER[g.weekday()], "tahsilat": th, "odeme": od, "net": th - od, "kumulatif": kum})
        if en_dusuk is None or kum < en_dusuk["kumulatif"]:
            en_dusuk = gunluk[-1]
    if kasa is not None and en_dusuk["kumulatif"] < 0:
        ilk = next(g for g in gunluk if g["kumulatif"] < 0) if baslangic >= 0 else None
        uy.append(uyari("Yüksek", "Nakit açığı", "Dönem", (f"Vadesi geçmiş ödemeler ({o['gecmis']:,.2f}) başlangıç kasasını ({kasa:,.2f}) aşıyor. " if baslangic < 0 else f"Kümülatif nakit {ilk['tarih']:%d.%m.%Y} günü eksiye düşüyor. ")
                        + f"En düşük nokta {en_dusuk['tarih']:%d.%m.%Y}: {en_dusuk['kumulatif']:,.2f}. Tahsilatı öne çekin, ödemeyi erteleyin veya kredi kullanımını planlayın."))
    for k in t["satirlar"]:
        if k["gecikme"] > 30:
            uy.append(uyari("Orta", "30 günü aşan gecikmiş alacak", f"{k['cari']} · {k['belge']}".strip(" ·"), f"{k['tutar']:,.2f} tutarındaki alacak {k['gecikme']} gündür gecikmede (vade {k['vade']:%d.%m.%Y})."))
    for k in o["satirlar"]:
        if k["gecikme"]:
            uy.append(uyari("Orta", "Vadesi geçmiş ödeme", f"{k['cari']} · {k['belge']}".strip(" ·"), f"{k['tutar']:,.2f} tutarındaki ödeme {k['gecikme']} gündür gecikmede (vade {k['vade']:%d.%m.%Y}). Gecikme faizi ve tedarik riski doğabilir."))
    for g in gunluk:
        if g["gun"] in ("Cmt", "Paz") and (g["tahsilat"] or g["odeme"]):
            uy.append(uyari("Bilgi", "Hafta sonuna denk gelen vade", f"{g['tarih']:%d.%m.%Y} {g['gun']}", "Bankalar kapalıysa işlem izleyen iş gününe kayar; günlük nakit etkisini buna göre değerlendirin."))
    return {"tahsilat": t, "odeme": o, "gunluk": gunluk, "bas": bas, "son": son, "kasa": kasa, "baslangic": baslangic, "net": t["donem"] - o["donem"], "en_dusuk": en_dusuk, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    for ad, veri, kisi in (("Tahsilat Listesi", r["tahsilat"], "Müşteri"), ("Ödeme Listesi", r["odeme"], "Tedarikçi")):
        ws = wb.active if ad == "Tahsilat Listesi" else wb.create_sheet(ad)
        ws.title = ad
        _baslik(ws, ["Durum", "Vade", "Gün", kisi, "Belge No", "Tutar", "Gecikme (gün)", "Açıklama", "Sonuç / Not (siz doldurun)"], (14, 12, 6, 32, 16, 16, 10, 30, 30))
        for k in veri["satirlar"]:
            _satir(ws, [k["durum"], k["vade"], GUNLER[k["vade"].weekday()], k["cari"], k["belge"], k["tutar"], k["gecikme"] or None, k["aciklama"], ""], para_s=(6,), tarih_s=(2,), dolgu={1: KIRMIZI if k["durum"] == "Vadesi geçmiş" else None})
        ws.append([])
        for etiket, tutar in (("Vadesi geçmiş", veri["gecmis"]), ("Bu dönem", veri["donem"]), ("Liste toplamı", veri["toplam"]), ("İleri vadeli (listede yok)", veri["ileri"])):
            n = _satir(ws, [etiket, None, None, None, None, tutar], para_s=(6,))
            ws.cell(n, 1).font = ws.cell(n, 6).font = Font(bold=True)
    ws = wb.create_sheet("Günlük Özet")
    _baslik(ws, ["Tarih", "Gün", "Tahsilat", "Ödeme", "Net", "Kümülatif Nakit"], (12, 6, 16, 16, 16, 18))
    _satir(ws, ["Başlangıç", "", None, r["odeme"]["gecmis"] or None, None, r["baslangic"]], para_s=(4, 6))
    for g in r["gunluk"]:
        _satir(ws, [g["tarih"], g["gun"], g["tahsilat"] or None, g["odeme"] or None, g["net"], g["kumulatif"]], para_s=(3, 4, 5, 6), tarih_s=(1,), dolgu={6: KIRMIZI if g["kumulatif"] < 0 else None, 2: MAVI if g["gun"] in ("Cmt", "Paz") else None})
    ws.append([])
    ws.append([f"Dönem {r['bas']:%d.%m.%Y} – {r['son']:%d.%m.%Y}. Başlangıç = kasa ({(r['kasa'] or Z):,.2f}) − vadesi geçmiş ödemeler. Vadesi geçmiş alacaklar kümülatife dahil değildir (ihtiyatlı). Vade günü tahsilat günü değildir."])
    ws = wb.create_sheet("Cari Özeti")
    _baslik(ws, ["Tür", "Cari", "Kalem", "Vadesi Geçmiş", "Bu Dönem", "Toplam", "En Eski Vade"], (12, 32, 8, 16, 16, 16, 13))
    for tur, veri in (("Tahsilat", r["tahsilat"]), ("Ödeme", r["odeme"])):
        for c in veri["cariler"]:
            _satir(ws, [tur, c["cari"], c["kalem"], c["gecmis"] or None, c["donem"] or None, c["toplam"], c["en_eski"]], para_s=(4, 5, 6), tarih_s=(7,))
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(alacak_yol: Path | None, borc_yol: Path | None, cikti: Path, bas: date, gun: int = 7, kasa: Decimal | None = None) -> dict:
    alacaklar, borclar, uy = [], [], []
    if alacak_yol:
        alacaklar, u = oku(alacak_yol, "Alacak")
        uy += u
    if borc_yol:
        borclar, u = oku(borc_yol, "Borç")
        uy += u
    r = hazirla(alacaklar, borclar, bas, gun, kasa)
    uy += r["uyarilar"]
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Açık alacak ve borçlardan haftalık tahsilat ve ödeme listesi çıkarır.")
    p.add_argument("--alacaklar", type=Path, help="Cari, Belge No, Vade, Tutar, Açıklama")
    p.add_argument("--borclar", type=Path, help="Cari, Belge No, Vade, Tutar, Açıklama")
    p.add_argument("--tarih", help="Dönem başlangıcı (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--gun", type=int, default=7, help="Dönem uzunluğu, gün (varsayılan 7)")
    p.add_argument("--kasa", help="Dönem başı kullanılabilir nakit (kasa ve banka); verilirse nakit açığı uyarısı üretilir")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "tahsilat_odeme_listesi.xlsx")
    a = p.parse_args(argv)
    ornek = a.alacaklar is None and a.borclar is None
    alacak, borc = (ORNEK / "alacaklar.csv", ORNEK / "borclar.csv") if ornek else (a.alacaklar, a.borclar)
    bas = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    kasa = para(a.kasa) if a.kasa else Decimal(10000) if ornek else None
    if bas is None or not 1 <= a.gun <= 92 or (a.kasa and kasa is None):
        print("[X] --tarih GG.AA.YYYY biçiminde, --gun 1–92 arasında, --kasa sayı olmalı")
        return 2
    for y in (alacak, borc):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(alacak, borc, a.cikti, bas, a.gun, kasa)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t, o = r["tahsilat"], r["odeme"]
    print(f"[OK] {r['bas']:%d.%m.%Y} – {r['son']:%d.%m.%Y} · tahsilat: {t['donem']:,.2f} bu dönem + {t['gecmis']:,.2f} gecikmiş · ödeme: {o['donem']:,.2f} bu dönem + {o['gecmis']:,.2f} gecikmiş · dönem neti {r['net']:,.2f}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
