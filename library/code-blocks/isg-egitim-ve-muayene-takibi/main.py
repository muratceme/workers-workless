"""
İSG Eğitim ve Muayene Takibi — Workers / Workless kod bloğu
İş Sağlığı ve Güvenliği › İş Güvenliği Uzmanı

Çalışanların tehlike sınıfına göre İSG eğitimi ve periyodik sağlık muayenesi tarihlerini takip eder:
  Eğitim (Çalışanların İSG Eğitimlerinin Usul ve Esasları Hakkında Yönetmelik, RG 02.04.2026 / 33212):
    - İşe başlama eğitimi en az 2 saat, işe başlamadan önce (md. 7/5).
    - Temel eğitim en az 8 / 12 / 16 ders saati (az tehlikeli / tehlikeli / çok tehlikeli; md. 13/1); işe başladıktan
      sonra en geç 3 ay içinde tamamlanır (md. 8/1).
    - Tekrar: 3 / 2 / 1 yılda en az bir (md. 14/1); tekrar eğitimi en az 8 ders saati (md. 14/2).
    - 6 aydan fazla işten uzak kalan çalışana işe başlatılmadan önce bilgi yenileme eğitimi (md. 18/1); iş kazası veya
      meslek hastalığı sonrası işe dönüşte ilave eğitim (md. 19/1).
  Muayene (İşyeri Hekimi ve Diğer Sağlık Personelinin Görev, Yetki, Sorumluluk ve Eğitimleri Hakkında Yönetmelik):
    - İşe giriş muayenesi işe başlamadan önce; periyodik muayene en geç 5 / 3 / 1 yılda bir. İşyeri hekiminin
      belirlediği daha kısa süre "Özel Muayene Periyodu (ay)" sütunuyla girilir.
  İlkyardım (İlkyardım Yönetmeliği): her 20 / 15 / 10 çalışana bir ilkyardımcı; ilkyardımcı belgesi 3 yıl geçerli.
  Eğitim döngüsü: saatler birikir; ilk döngü temel eğitim saatine, sonrakiler tekrar saatine ulaştığı gün tamamlanır.
Rapor: durum listesi, aksiyon listesi (tarih sırasıyla), aylık plan, birim özeti, ilkyardımcı yeterliliği, uyarılar.
İnternete bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: 26 aktif çalışan, 2 işyeri
    python main.py --calisanlar c.xlsx --egitimler e.xlsx --muayeneler m.xlsx --tehlike tehlikeli --bugun 09.10.2026
"""
from __future__ import annotations

import argparse
import calendar
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
AYLAR = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]
SINIFLAR = {"az": "Az tehlikeli", "tehlikeli": "Tehlikeli", "cok": "Çok tehlikeli"}
TEMEL_SAAT = {"az": 8, "tehlikeli": 12, "cok": 16}
TEKRAR_YIL = {"az": 3, "tehlikeli": 2, "cok": 1}
MUAYENE_YIL = {"az": 5, "tehlikeli": 3, "cok": 1}
ILKYARDIM_ORAN = {"az": 20, "tehlikeli": 15, "cok": 10}
TEKRAR_SAAT, ISE_BASLAMA_SAAT, TEMEL_AY, ILKYARDIM_YIL = 8, 2, 3, 3

CALISAN_SUTUNLARI = {"sicil": ("sicil", "sicil no", "personel no"), "ad": ("ad soyad", "adi soyadi", "calisan"), "birim": ("birim", "bolum", "departman"),
                     "gorev": ("gorev", "unvan", "pozisyon"), "isyeri": ("isyeri", "lokasyon", "tesis"), "giris": ("ise giris", "ise giris tarihi", "giris tarihi"),
                     "cikis": ("isten cikis", "cikis tarihi", "ayrilis tarihi"), "sinif": ("tehlike sinifi",), "ozel": ("ozel muayene periyodu", "ozel periyot", "ozel muayene periyodu ay"),
                     "uzun": ("uzun ayrilik donusu", "uzun ayrilik donus tarihi"), "kaza": ("is kazasi donusu", "is kazasi donus tarihi", "meslek hastaligi donusu")}
EGITIM_SUTUNLARI = {"sicil": ("sicil", "sicil no", "personel no"), "tarih": ("tarih", "egitim tarihi"), "tur": ("egitim turu", "tur", "egitim"),
                    "saat": ("sure", "saat", "sure saat", "ders saati", "sure ders saati", "sure saat"), "egitici": ("egitici", "veren")}
MUAYENE_SUTUNLARI = {"sicil": ("sicil", "sicil no", "personel no"), "tarih": ("tarih", "muayene tarihi"), "tur": ("muayene turu", "tur"), "sonuc": ("sonuc", "karar")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def sayi(x) -> Decimal | None:
    if x in (None, ""):
        return None
    try:
        return Decimal(str(x).replace(",", ".").strip())
    except InvalidOperation:
        return None


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def ay_ekle(d: date, n: int) -> date:
    y, m = divmod(d.month - 1 + n, 12)
    yil, ay = d.year + y, m + 1
    return date(yil, ay, min(d.day, calendar.monthrange(yil, ay)[1]))


def sinif_bul(x) -> str | None:
    k = katla(x)
    if not k:
        return None
    if k.startswith("cok"):
        return "cok"
    if k.startswith("az"):
        return "az"
    return "tehlikeli" if "tehlikeli" in k else None


def egitim_turu(x) -> str:
    k = katla(x)
    if "ilk yardim" in k or "ilkyardim" in k:
        return "ilkyardim"
    if "ise baslama" in k or "oryantasyon" in k:
        return "ise_baslama"
    if "bilgi yenileme" in k:
        return "bilgi_yenileme"
    if "ilave" in k or "is kazasi" in k or "meslek hastaligi" in k:
        return "ilave"
    if "tekrar" in k or "periyodik" in k or "yenileme" in k:
        return "tekrar"
    if "temel" in k or "isg" in k or "is sagligi" in k:
        return "temel"
    return "diger"


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Calisan:
    sicil: str
    ad: str
    birim: str
    gorev: str
    isyeri: str
    giris: date
    cikis: date | None
    sinif: str
    ozel_ay: int | None
    uzun_donus: date | None
    kaza_donus: date | None
    egitimler: list[tuple[date, str, Decimal]] = field(default_factory=list)
    muayeneler: list[tuple[date, str, str]] = field(default_factory=list)
    durum: dict = field(default_factory=dict)


def oku(c_yol: Path, e_yol: Path | None, m_yol: Path | None, varsayilan_sinif: str, bugun: date) -> tuple[list[Calisan], list[dict]]:
    uy, calisanlar = [], {}
    for r in kayitlar(c_yol, CALISAN_SUTUNLARI, ("sicil", "giris")):
        s = metin(r.get("sicil"))
        if not s:
            continue
        g = tarih(r.get("giris"))
        if not g:
            uy.append({"onem": "Yüksek", "tur": "İşe giriş tarihi yok", "kim": s, "aciklama": "Çalışan alınmadı; işe giriş tarihini girin"})
            continue
        cikis = tarih(r.get("cikis"))
        if cikis and cikis <= bugun:
            continue
        sn = sinif_bul(r.get("sinif"))
        if metin(r.get("sinif")) and not sn:
            uy.append({"onem": "Orta", "tur": "Tehlike sınıfı okunamadı", "kim": s, "aciklama": f"'{metin(r.get('sinif'))}'; işyeri sınıfı kullanıldı"})
        oz = sayi(r.get("ozel"))
        calisanlar[s] = Calisan(s, metin(r.get("ad")), metin(r.get("birim")) or "(belirtilmemiş)", metin(r.get("gorev")), metin(r.get("isyeri")) or "Merkez", g, cikis,
                                sn or varsayilan_sinif, int(oz) if oz else None, tarih(r.get("uzun")), tarih(r.get("kaza")))
    bilinmeyen = defaultdict(int)
    if e_yol:
        for r in kayitlar(e_yol, EGITIM_SUTUNLARI, ("sicil", "tarih", "tur")):
            c = calisanlar.get(metin(r.get("sicil")))
            t = tarih(r.get("tarih"))
            if not c or not t:
                if metin(r.get("sicil")) and not c:
                    bilinmeyen["eğitim"] += 1
                continue
            c.egitimler.append((t, egitim_turu(r.get("tur")), sayi(r.get("saat")) or Decimal(0)))
    if m_yol:
        for r in kayitlar(m_yol, MUAYENE_SUTUNLARI, ("sicil", "tarih")):
            c = calisanlar.get(metin(r.get("sicil")))
            t = tarih(r.get("tarih"))
            if not c or not t:
                if metin(r.get("sicil")) and not c:
                    bilinmeyen["muayene"] += 1
                continue
            c.muayeneler.append((t, katla(r.get("tur")), metin(r.get("sonuc"))))
    for k, n in bilinmeyen.items():
        uy.append({"onem": "Bilgi", "tur": "Eşleşmeyen kayıt", "kim": k, "aciklama": f"{n} {k} kaydı aktif çalışan listesinde olmayan sicile ait (ayrılmış olabilir)"})
    return list(calisanlar.values()), uy


# ----------------------------------------------------------------------------
# Değerlendirme
# ----------------------------------------------------------------------------

def egitim_donguleri(c: Calisan) -> list[date]:
    """Temel + tekrar eğitim saatlerini biriktirir; ilk döngü temel saatine, sonrakiler tekrar saatine ulaştığında tamamlanır."""
    tamam, birikim = [], Decimal(0)
    for t, tur, saat in sorted(e for e in c.egitimler if e[1] in ("temel", "tekrar")):
        birikim += saat
        gerek = TEMEL_SAAT[c.sinif] if not tamam else TEKRAR_SAAT
        if birikim >= gerek:
            tamam.append(t)
            birikim = Decimal(0)
    return tamam


def durum_yaz(son: date | None, bugun: date, yaklasan: int) -> str:
    if son is None:
        return "—"
    kalan = (son - bugun).days
    return "Gecikmiş" if kalan < 0 else "Yaklaşan" if kalan <= yaklasan else "Uygun"


def degerlendir(calisanlar: list[Calisan], bugun: date, yaklasan: int = 60) -> dict:
    uy, aksiyon = [], []

    def ekle(c, konu, son, onem, aciklama):
        aksiyon.append({"tarih": son, "sicil": c.sicil, "ad": c.ad, "birim": c.birim, "konu": konu, "onem": onem, "aciklama": aciklama})
        uy.append({"onem": onem, "tur": konu, "kim": f"{c.sicil} {c.ad}".strip(), "aciklama": aciklama})

    for c in calisanlar:
        d = c.durum
        # İşe başlama eğitimi
        ib = [e for e in c.egitimler if e[1] == "ise_baslama"]
        ib_saat = sum((e[2] for e in ib if e[0] <= c.giris), Decimal(0))
        d["ise_baslama"] = "Uygun" if ib_saat >= ISE_BASLAMA_SAAT else ("Geç / eksik" if ib else "Yok")
        if d["ise_baslama"] != "Uygun" and c.giris >= date(2026, 4, 2):
            ekle(c, "İşe başlama eğitimi", c.giris, "Orta",
                 f"İşe başlamadan önce en az {ISE_BASLAMA_SAAT} saat olmalı (md. 7/5); işe girişe kadar {ib_saat:g} saat kayıtlı")
        # Temel / tekrar eğitim
        dongu = egitim_donguleri(c)
        d["egitim_son"] = dongu[-1] if dongu else None
        if dongu:
            d["egitim_sonraki"] = ay_ekle(dongu[-1], 12 * TEKRAR_YIL[c.sinif])
            d["egitim_durum"] = durum_yaz(d["egitim_sonraki"], bugun, yaklasan)
        else:
            d["egitim_sonraki"] = ay_ekle(c.giris, TEMEL_AY)
            birikim = sum((e[2] for e in c.egitimler if e[1] in ("temel", "tekrar")), Decimal(0))
            d["egitim_durum"] = "Temel eğitim gecikmiş" if d["egitim_sonraki"] < bugun else "Temel eğitim devam ediyor"
            d["egitim_birikim"] = birikim
        if d["egitim_durum"] in ("Gecikmiş", "Temel eğitim gecikmiş"):
            ekle(c, "İSG eğitimi", d["egitim_sonraki"], "Yüksek",
                 f"Son gün {d['egitim_sonraki']:%d.%m.%Y} geçti; " + (f"son tamamlanan döngü {d['egitim_son']:%d.%m.%Y} ({TEKRAR_YIL[c.sinif]} yılda bir, md. 14/1)"
                                                                    if d["egitim_son"] else f"temel eğitim {TEMEL_SAAT[c.sinif]} saate ulaşmadı "
                                                                                             f"({d['egitim_birikim']:g} saat; md. 8/1, 13/1)"))
        elif d["egitim_durum"] in ("Yaklaşan", "Temel eğitim devam ediyor"):
            ekle(c, "İSG eğitimi", d["egitim_sonraki"], "Orta" if d["egitim_durum"] == "Yaklaşan" else "Bilgi",
                 f"Son gün {d['egitim_sonraki']:%d.%m.%Y}; " + ("tekrar eğitimi en az 8 ders saati" if d["egitim_son"] else
                                                              f"temel eğitim {TEMEL_SAAT[c.sinif]} saatin {d.get('egitim_birikim', 0):g} saati tamam"))
        # Bilgi yenileme / ilave eğitim
        if c.uzun_donus:
            ok = any(e[1] == "bilgi_yenileme" and c.uzun_donus - timedelta(days=30) <= e[0] <= c.uzun_donus for e in c.egitimler)
            d["bilgi_yenileme"] = "Uygun" if ok else "Eksik"
            if not ok:
                ekle(c, "Bilgi yenileme eğitimi", c.uzun_donus, "Yüksek",
                     f"6 aydan uzun ayrılık sonrası {c.uzun_donus:%d.%m.%Y} dönüşünden önce bilgi yenileme eğitimi kaydı yok (md. 18/1)")
        if c.kaza_donus:
            ok = any(e[1] == "ilave" and c.kaza_donus - timedelta(days=30) <= e[0] <= c.kaza_donus + timedelta(days=7) for e in c.egitimler)
            d["ilave"] = "Uygun" if ok else "Eksik"
            if not ok:
                ekle(c, "İlave eğitim (iş kazası / meslek hastalığı)", c.kaza_donus, "Yüksek",
                     f"{c.kaza_donus:%d.%m.%Y} işe dönüşünde ilave eğitim kaydı yok (md. 19/1)")
        # Muayene
        mu = sorted(c.muayeneler)
        giris_m = [m for m in mu if "giris" in m[1] or "ise baslama" in m[1]]
        d["ise_giris_muayene"] = "Uygun" if any(m[0] <= c.giris for m in giris_m) else ("Geç" if giris_m else "Yok")
        if d["ise_giris_muayene"] != "Uygun" and (bugun - c.giris).days <= 365:
            ekle(c, "İşe giriş muayenesi", c.giris, "Yüksek" if d["ise_giris_muayene"] == "Yok" else "Orta",
                 "İşe giriş muayenesi işe başlamadan önce yapılmalı" + (f"; kayıt {giris_m[0][0]:%d.%m.%Y}" if giris_m else "; kayıt yok"))
        periyot_ay = min(12 * MUAYENE_YIL[c.sinif], c.ozel_ay) if c.ozel_ay else 12 * MUAYENE_YIL[c.sinif]
        d["periyot_ay"] = periyot_ay
        d["muayene_son"] = mu[-1][0] if mu else None
        d["muayene_sonraki"] = ay_ekle(mu[-1][0], periyot_ay) if mu else c.giris
        d["muayene_durum"] = durum_yaz(d["muayene_sonraki"], bugun, yaklasan)
        d["kisitli"] = mu[-1][2] if mu and katla(mu[-1][2]) not in ("", "calisabilir", "uygun", "calismaya elverislidir") else ""
        if not mu:
            ekle(c, "Muayene kaydı yok", c.giris, "Yüksek", "Hiç sağlık muayenesi kaydı yok; işe giriş / periyodik muayene belgesini kontrol edin")
        elif d["muayene_durum"] == "Gecikmiş":
            ekle(c, "Periyodik muayene", d["muayene_sonraki"], "Yüksek",
                 f"Son muayene {mu[-1][0]:%d.%m.%Y}; periyot {periyot_ay} ay → {d['muayene_sonraki']:%d.%m.%Y} geçti")
        elif d["muayene_durum"] == "Yaklaşan":
            ekle(c, "Periyodik muayene", d["muayene_sonraki"], "Orta", f"Son gün {d['muayene_sonraki']:%d.%m.%Y} (periyot {periyot_ay} ay)")
        if d["kisitli"]:
            uy.append({"onem": "Bilgi", "tur": "Kısıtlı / şartlı rapor", "kim": f"{c.sicil} {c.ad}".strip(),
                       "aciklama": f"Son muayene sonucu: {d['kisitli']}; görevlendirmeyi işyeri hekimiyle kontrol edin"})
        if c.kaza_donus:
            d["donus_muayene"] = "Var" if any(m[0] >= c.kaza_donus - timedelta(days=30) for m in mu) else "Yok"
            if d["donus_muayene"] == "Yok":
                uy.append({"onem": "Bilgi", "tur": "İşe dönüş muayenesi", "kim": f"{c.sicil} {c.ad}".strip(),
                           "aciklama": "İş kazası / meslek hastalığı sonrası dönüşte muayene kaydı yok; çalışanın talebi ve işyeri hekiminin değerlendirmesi için kontrol edin"})
        # İlkyardım belgesi
        ilk = [e[0] for e in c.egitimler if e[1] == "ilkyardim"]
        d["ilkyardim_bitis"] = ay_ekle(max(ilk), 12 * ILKYARDIM_YIL) if ilk else None
        d["ilkyardimci"] = bool(ilk) and d["ilkyardim_bitis"] >= bugun
        if ilk and d["ilkyardim_bitis"] < bugun + timedelta(days=yaklasan):
            ekle(c, "İlkyardımcı belgesi", d["ilkyardim_bitis"], "Orta" if d["ilkyardim_bitis"] >= bugun else "Yüksek",
                 f"İlkyardımcı belgesi {d['ilkyardim_bitis']:%d.%m.%Y} tarihinde " + ("bitiyor" if d["ilkyardim_bitis"] >= bugun else "bitti") + "; güncelleme eğitimi planlayın")

    # İlkyardımcı yeterliliği (işyeri bazında, en yüksek tehlike sınıfı oranıyla)
    ilkyardim = []
    for isyeri in sorted({c.isyeri for c in calisanlar}):
        lst = [c for c in calisanlar if c.isyeri == isyeri]
        oran = min(ILKYARDIM_ORAN[c.sinif] for c in lst)
        gerek = math.ceil(len(lst) / oran)
        mevcut = sum(c.durum["ilkyardimci"] for c in lst)
        ilkyardim.append({"isyeri": isyeri, "calisan": len(lst), "oran": oran, "gerek": gerek, "mevcut": mevcut})
        if mevcut < gerek:
            uy.append({"onem": "Yüksek", "tur": "İlkyardımcı yetersiz", "kim": isyeri,
                       "aciklama": f"{len(lst)} çalışan, her {oran} çalışana bir → en az {gerek} ilkyardımcı; geçerli belgeli {mevcut}"})
    birim = defaultdict(lambda: defaultdict(int))
    for c in calisanlar:
        b = birim[c.birim]
        b["calisan"] += 1
        b["egitim_gecikmis"] += c.durum["egitim_durum"] in ("Gecikmiş", "Temel eğitim gecikmiş")
        b["egitim_yaklasan"] += c.durum["egitim_durum"] == "Yaklaşan"
        b["muayene_gecikmis"] += c.durum["muayene_durum"] == "Gecikmiş"
        b["muayene_yaklasan"] += c.durum["muayene_durum"] == "Yaklaşan"
    plan = defaultdict(lambda: defaultdict(list))
    for c in calisanlar:
        for konu, t in (("İSG eğitimi", c.durum["egitim_sonraki"]), ("Periyodik muayene", c.durum["muayene_sonraki"])):
            if t and t <= ay_ekle(bugun, 6):
                ay = (t.year, t.month) if t >= bugun else (bugun.year, bugun.month)
                plan[ay][konu].append(c)
    aksiyon.sort(key=lambda a: (a["tarih"], a["sicil"]))
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    uy.sort(key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))
    return {"calisanlar": calisanlar, "aksiyon": aksiyon, "uyarilar": uy, "ilkyardim": ilkyardim, "birim": birim, "plan": dict(sorted(plan.items())), "bugun": bugun}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
DURUM_RENK = {"Gecikmiş": "FDE2E1", "Temel eğitim gecikmiş": "FDE2E1", "Yaklaşan": "FFF4CE", "Temel eğitim devam ediyor": "E8F0FE", "Uygun": "E3F4E1",
              "Yok": "FDE2E1", "Eksik": "FDE2E1", "Geç": "FFF4CE", "Geç / eksik": "FFF4CE"}
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, basliklar, genislik=()):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "C2"


def rapor_yaz(cikti: Path, s: dict) -> None:
    wb = Workbook()
    du = wb.active
    du.title = "Durum"
    _baslik(du, ["Sicil", "Ad Soyad", "Birim", "Görev", "İşyeri", "Tehlike Sınıfı", "İşe Giriş", "İşe Başlama Eğt.", "Son Eğitim Döngüsü", "Sonraki Eğitim",
                 "Eğitim Durumu", "İşe Giriş Muayenesi", "Son Muayene", "Periyot (ay)", "Sonraki Muayene", "Muayene Durumu", "İlkyardımcı Belgesi Bitiş",
                 "Bilgi Yenileme", "İlave Eğitim", "Son Muayene Sonucu"], (9, 22, 14, 18, 10, 12, 11, 11, 11, 11, 16, 11, 11, 8, 11, 12, 12, 10, 10, 24))
    for c in sorted(s["calisanlar"], key=lambda c: (c.birim, c.sicil)):
        d = c.durum
        du.append([c.sicil, c.ad, c.birim, c.gorev, c.isyeri, SINIFLAR[c.sinif], c.giris, d["ise_baslama"], d["egitim_son"], d["egitim_sonraki"], d["egitim_durum"],
                   d["ise_giris_muayene"], d["muayene_son"], d["periyot_ay"], d["muayene_sonraki"], d["muayene_durum"], d["ilkyardim_bitis"],
                   d.get("bilgi_yenileme", ""), d.get("ilave", ""), d["kisitli"]])
        for j in (7, 9, 10, 13, 15, 17):
            du.cell(du.max_row, j).number_format = "DD.MM.YYYY"
        for j in (8, 11, 12, 16, 18, 19):
            renk = DURUM_RENK.get(du.cell(du.max_row, j).value)
            if renk:
                du.cell(du.max_row, j).fill = PatternFill("solid", fgColor=renk)
    du.auto_filter.ref = f"A1:T{du.max_row}"

    ak = wb.create_sheet("Aksiyon Listesi")
    _baslik(ak, ["Son Tarih", "Kalan Gün", "Önem", "Konu", "Sicil", "Ad Soyad", "Birim", "Açıklama", "Planlanan Tarih", "Yapıldı"], (11, 8, 8, 30, 9, 22, 14, 80, 12, 9))
    ak.freeze_panes = "A2"
    for a in s["aksiyon"]:
        ak.append([a["tarih"], (a["tarih"] - s["bugun"]).days, a["onem"], a["konu"], a["sicil"], a["ad"], a["birim"], a["aciklama"], None, ""])
        ak.cell(ak.max_row, 1).number_format = ak.cell(ak.max_row, 9).number_format = "DD.MM.YYYY"
        ak.cell(ak.max_row, 3).fill = PatternFill("solid", fgColor=RENK[a["onem"]])
        ak.cell(ak.max_row, 8).alignment = UST
        for j in (9, 10):
            ak.cell(ak.max_row, j).fill = PatternFill("solid", fgColor="FFF4CE")

    pl = wb.create_sheet("Aylık Plan")
    _baslik(pl, ["Ay", "Konu", "Kişi", "Çalışanlar (sicil – ad)"], (14, 18, 6, 110))
    pl.freeze_panes = "A2"
    for (y, m), konular in s["plan"].items():
        for konu, lst in sorted(konular.items()):
            pl.append([f"{AYLAR[m - 1]} {y}" + (" (gecikmiş dahil)" if (y, m) == (s["bugun"].year, s["bugun"].month) else ""), konu, len(lst),
                       "; ".join(f"{c.sicil} {c.ad}" for c in sorted(lst, key=lambda c: c.sicil))])
            pl.cell(pl.max_row, 4).alignment = UST

    bo = wb.create_sheet("Birim Özeti")
    _baslik(bo, ["Birim", "Çalışan", "Eğitim Gecikmiş", "Eğitim Yaklaşan", "Muayene Gecikmiş", "Muayene Yaklaşan"], (18, 9, 10, 10, 10, 10))
    for b, v in sorted(s["birim"].items()):
        bo.append([b, v["calisan"], v["egitim_gecikmis"], v["egitim_yaklasan"], v["muayene_gecikmis"], v["muayene_yaklasan"]])
        for j in (3, 5):
            if bo.cell(bo.max_row, j).value:
                bo.cell(bo.max_row, j).fill = PatternFill("solid", fgColor="FDE2E1")

    il = wb.create_sheet("İlkyardımcı")
    _baslik(il, ["İşyeri", "Çalışan", "Her Kaç Çalışana Bir", "Gereken", "Geçerli Belgeli", "Fark"], (16, 9, 12, 9, 10, 8))
    for x in s["ilkyardim"]:
        il.append([x["isyeri"], x["calisan"], x["oran"], x["gerek"], x["mevcut"], x["mevcut"] - x["gerek"]])
        il.cell(il.max_row, 6).fill = PatternFill("solid", fgColor="E3F4E1" if x["mevcut"] >= x["gerek"] else "FDE2E1")
    il.append([])
    il.append(["Oran, işyerindeki en yüksek tehlike sınıfına göre alınır (20 / 15 / 10). İlkyardımcı belgesi 3 yıl geçerli sayılır."])

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Kim", "Açıklama"], (9, 30, 26, 100))
    uy.freeze_panes = "A2"
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["kim"], u["aciklama"]])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(c_yol: Path, cikti: Path, bugun: date, e_yol: Path | None = None, m_yol: Path | None = None, tehlike: str = "tehlikeli", yaklasan: int = 60) -> dict:
    calisanlar, uy = oku(c_yol, e_yol, m_yol, tehlike, bugun)
    if not calisanlar:
        raise ValueError(f"{c_yol.name}: aktif çalışan bulunamadı")
    s = degerlendir(calisanlar, bugun, yaklasan)
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    s["uyarilar"] = sorted(uy + s["uyarilar"], key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))
    rapor_yaz(cikti, s)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Tehlike sınıfına göre İSG eğitimi ve periyodik sağlık muayenesi tarihlerini takip eder.")
    p.add_argument("--calisanlar", type=Path, default=ORNEK / "calisanlar.csv",
                   help="Sicil, Ad Soyad, Birim, Görev, İşyeri, İşe Giriş, İşten Çıkış, Tehlike Sınıfı, Özel Muayene Periyodu (ay), Uzun Ayrılık Dönüşü, "
                        "İş Kazası Dönüşü")
    p.add_argument("--egitimler", type=Path, help="Sicil, Tarih, Eğitim Türü (İşe başlama / Temel / Tekrar / Bilgi yenileme / İlave / İlkyardım), Süre (ders saati)")
    p.add_argument("--muayeneler", type=Path, help="Sicil, Tarih, Muayene Türü (İşe giriş / Periyodik / İşe dönüş), Sonuç")
    p.add_argument("--tehlike", default="tehlikeli", help="İşyeri tehlike sınıfı: az, tehlikeli, cok (çalışan satırında yoksa; varsayılan tehlikeli)")
    p.add_argument("--yaklasan", type=int, default=60, help="Son tarihe bu kadar gün kala 'yaklaşan' (varsayılan 60)")
    p.add_argument("--bugun", help="Rapor tarihi GG.AA.YYYY (örnek veride 09.10.2026)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "isg_takip.xlsx")
    a = p.parse_args(argv)
    ornek = a.calisanlar == ORNEK / "calisanlar.csv"
    e = a.egitimler or (ORNEK / "egitimler.csv" if ornek else None)
    m = a.muayeneler or (ORNEK / "muayeneler.csv" if ornek else None)
    for y in (a.calisanlar, e, m):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    sinif = sinif_bul(a.tehlike.replace("cok", "çok"))
    bugun = tarih(a.bugun) if a.bugun else (date(2026, 10, 9) if ornek else date.today())
    if not sinif or not bugun:
        print("[X] --tehlike az / tehlikeli / cok ve --bugun GG.AA.YYYY olmalı")
        return 2
    try:
        s = calistir(a.calisanlar, a.cikti, bugun, e, m, sinif, a.yaklasan)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    cl = s["calisanlar"]
    print(f"[OK] {len(cl)} aktif çalışan · eğitimi gecikmiş {sum(c.durum['egitim_durum'] in ('Gecikmiş', 'Temel eğitim gecikmiş') for c in cl)} · "
          f"muayenesi gecikmiş {sum(c.durum['muayene_durum'] == 'Gecikmiş' for c in cl)} · {len(s['aksiyon'])} aksiyon")
    for u in s["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['kim']} · {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
