import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 9)


def calisan(sinif="tehlikeli", giris=date(2020, 1, 6), egitimler=(), muayeneler=(), **kw):
    c = main.Calisan("X1", "Test", "Üretim", "Operatör", "Merkez", giris, None, sinif, kw.get("ozel"), kw.get("uzun"), kw.get("kaza"))
    c.egitimler = [(t, tur, D(s)) for t, tur, s in egitimler]
    c.muayeneler = [(t, tur, sonuc) for t, tur, sonuc in muayeneler]
    return c


class KuralTesti(unittest.TestCase):
    def test_egitim_dongusu(self):
        c = calisan(egitimler=[(date(2020, 1, 20), "temel", 8), (date(2020, 2, 3), "temel", 4), (date(2022, 1, 10), "tekrar", 4),
                               (date(2022, 1, 17), "tekrar", 4), (date(2024, 1, 5), "tekrar", 6)])
        self.assertEqual(main.egitim_donguleri(c), [date(2020, 2, 3), date(2022, 1, 17)])     # 12 saat temel; tekrar 8 saat; son 6 saat yetmez
        c = calisan("cok", egitimler=[(date(2020, 1, 20), "temel", 12)])
        self.assertEqual(main.egitim_donguleri(c), [])                                         # çok tehlikelide 16 saat gerekir

    def test_periyotlar(self):
        s = main.degerlendir([calisan("az", egitimler=[(date(2020, 1, 6), "ise_baslama", 2), (date(2023, 11, 1), "temel", 8)],
                                      muayeneler=[(date(2020, 1, 2), "ise giris", "Çalışabilir"), (date(2022, 5, 1), "periyodik", "")])], BUGUN)
        d = s["calisanlar"][0].durum
        self.assertEqual((d["egitim_sonraki"], d["egitim_durum"]), (date(2026, 11, 1), "Yaklaşan"))      # az tehlikeli: 3 yıl
        self.assertEqual((d["muayene_sonraki"], d["muayene_durum"]), (date(2027, 5, 1), "Uygun"))        # 5 yıl
        s = main.degerlendir([calisan("tehlikeli", muayeneler=[(date(2025, 3, 1), "periyodik", "")], ozel=12)], BUGUN)
        self.assertEqual(s["calisanlar"][0].durum["muayene_sonraki"], date(2026, 3, 1))                   # özel periyot 12 ay

    def test_yeni_giris_ve_donusler(self):
        c = calisan(giris=date(2026, 8, 1), egitimler=[(date(2026, 8, 1), "ise_baslama", 2), (date(2026, 8, 10), "temel", 8)],
                    muayeneler=[(date(2026, 7, 28), "ise giris", "")])
        d = main.degerlendir([c], BUGUN)["calisanlar"][0].durum
        self.assertEqual((d["ise_baslama"], d["egitim_durum"], d["egitim_sonraki"], d["ise_giris_muayene"]),
                         ("Uygun", "Temel eğitim devam ediyor", date(2026, 11, 1), "Uygun"))
        c = calisan(uzun=date(2026, 9, 1), kaza=date(2026, 9, 20), egitimler=[(date(2026, 8, 28), "bilgi_yenileme", 2)])
        d = main.degerlendir([c], BUGUN)["calisanlar"][0].durum
        self.assertEqual((d["bilgi_yenileme"], d["ilave"]), ("Uygun", "Eksik"))

    def test_tur_ve_sinif(self):
        self.assertEqual(main.egitim_turu("Temel İSG eğitimi"), "temel")
        self.assertEqual(main.egitim_turu("İlk Yardım Güncelleme"), "ilkyardim")
        self.assertEqual(main.egitim_turu("Tekrar eğitimi"), "tekrar")
        self.assertEqual(main.egitim_turu("İlave eğitim (iş kazası sonrası)"), "ilave")
        self.assertEqual(main.sinif_bul("Çok Tehlikeli"), "cok")
        self.assertEqual(main.sinif_bul("az tehlikeli"), "az")
        self.assertIsNone(main.sinif_bul("orta"))


class OrnekVeriTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "i.xlsx"
            s = main.calistir(ORNEK / "calisanlar.csv", cikti, BUGUN, ORNEK / "egitimler.csv", ORNEK / "muayeneler.csv")
            self.assertEqual(len(s["calisanlar"]), 26)                                       # S030 ayrıldı
            d = {c.sicil: c.durum for c in s["calisanlar"]}
            self.assertEqual(d["S005"]["egitim_durum"], "Yaklaşan")
            self.assertEqual(d["S007"]["egitim_son"], date(2023, 6, 10))                     # 2025'teki 6 saat tekrar yetmez
            self.assertEqual(d["S008"]["egitim_sonraki"], date(2026, 3, 10))                 # çok tehlikeli: 1 yıl
            self.assertEqual(d["S009"]["muayene_durum"], "Yaklaşan")
            turler = {(u["tur"], u["kim"].split()[0]) for u in s["uyarilar"]}
            self.assertTrue({("Bilgi yenileme eğitimi", "S012"), ("İlave eğitim (iş kazası / meslek hastalığı)", "S014"), ("Muayene kaydı yok", "S017"),
                             ("Periyodik muayene", "S015"), ("İşe başlama eğitimi", "S011"), ("İşe giriş muayenesi", "S011"),
                             ("İlkyardımcı yetersiz", "Kocaeli"), ("Kısıtlı / şartlı rapor", "S016"), ("İşe dönüş muayenesi", "S013"),
                             ("Eşleşmeyen kayıt", "eğitim")} <= turler)
            self.assertNotIn(("İlave eğitim (iş kazası / meslek hastalığı)", "S013"), turler)
            self.assertEqual({x["isyeri"]: x["gerek"] for x in s["ilkyardim"]}, {"İstanbul Ofis": 1, "Kocaeli Fabrika": 3})   # 5 / 20; 21 / 10
            self.assertEqual([a["tarih"] for a in s["aksiyon"]], sorted(a["tarih"] for a in s["aksiyon"]))
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Durum", "Aksiyon Listesi", "Aylık Plan", "Birim Özeti", "İlkyardımcı", "Uyarılar"])

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--tehlike", "orta"]), 2)
            self.assertEqual(main.main(["--calisanlar", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
