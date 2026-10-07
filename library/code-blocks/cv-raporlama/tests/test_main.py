import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

KLASOR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KLASOR))

import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = KLASOR / "ornek_veri"


class DeneyimTesti(unittest.TestCase):
    def test_cakisan_araliklar_birlestirilir(self):
        metin = "A Şirketi 2015 - 2019\nB Şirketi 2018 - 2020"
        self.assertEqual(main.deneyim_hesapla(metin, date(2026, 1, 1)), 5.0)

    def test_gunumuz_bugune_kadar_sayilir(self):
        self.assertEqual(main.deneyim_hesapla("03/2021 - Günümüz", date(2026, 6, 1)), 5.0)

    def test_tarih_yoksa_sifir(self):
        self.assertEqual(main.deneyim_hesapla("deneyim yok"), 0.0)


class YardimciTesti(unittest.TestCase):
    def test_turkce_kucuk_harf(self):
        self.assertEqual(main.kucuk("İŞE ALIM"), "işe alım")

    def test_kelime_siniri(self):
        ms = main.sade("Java geliştirici")
        self.assertTrue(main.kelime_var(ms, "Java"))
        self.assertFalse(main.kelime_var(ms, "JavaScript"))

    def test_telefon_normallestirme(self):
        self.assertEqual(main.telefon_bul("Tel: 0 (532) 111-22-33"), "+90 532 111 22 33")


class UctanUcaTest(unittest.TestCase):
    def test_ornek_veri_raporu(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "rapor.xlsx"
            adaylar = main.calistir(ORNEK / "cvler", cikti, KLASOR / "beceriler.txt", ORNEK / "ilan.txt")
            by = {a.dosya: a for a in adaylar}

            self.assertEqual(len(adaylar), 4)
            ayse = by["ayse_yilmaz.txt"]
            self.assertEqual(ayse.ad_soyad, "Ayşe Yılmaz")
            self.assertEqual(ayse.telefon, "+90 532 111 22 33")
            self.assertEqual(ayse.deneyim_yil, float(date.today().year - 2015))
            self.assertEqual(ayse.egitim, "Lisans")
            self.assertIn("İngilizce", ayse.diller)
            self.assertGreater(ayse.ilan_uyum, 80)

            mehmet = by["mehmet_kaya.docx"]
            self.assertEqual(mehmet.ad_soyad, "Mehmet Kaya")
            self.assertEqual(mehmet.egitim, "Yüksek Lisans")

            self.assertEqual(by["zeynep_demir.pdf"].eposta, "zeynep.demir@ornek-mail.com")
            self.assertTrue(by["bozuk_dosya.pdf"].uyarilar)

            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Adaylar", "Özet"])
            self.assertEqual(wb["Adaylar"].max_row, 5)
            self.assertEqual(wb["Adaylar"]["B2"].value, "Ayşe Yılmaz")  # ilan uyumuna göre sıralı


if __name__ == "__main__":
    unittest.main()
