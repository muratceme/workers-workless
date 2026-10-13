"""
Porsiyon Maliyeti Hesaplama — Workers / Workless kod bloğu
Yiyecek ve İçecek › Yiyecek İçecek Müdürü

Standart reçete, gramaj ve güncel alış fiyatlarından menü kalemi başına porsiyon maliyetini ve maliyet oranını hesaplar.
  Brüt miktar        = reçetedeki net miktar ÷ (1 − fire % ÷ 100)                      (ayıklama, temizleme, pişirme kaybı)
  Malzeme maliyeti   = brüt miktar (alış birimine çevrilmiş) × alış fiyatı              (g ↔ kg, ml ↔ cl ↔ l çevrilir; adet yalnız adetle eşleşir)
  Porsiyon maliyeti  = Σ malzeme maliyeti ÷ reçetenin porsiyon sayısı                   (verilmezse 1)
  Maliyet oranı %    = porsiyon maliyeti ÷ satış fiyatı (KDV hariç) × 100
  Brüt kâr           = satış fiyatı (KDV hariç) − porsiyon maliyeti
  Hedef fiyat        = porsiyon maliyeti ÷ (--hedef % ÷ 100), KDV dahil hâliyle          (varsayılan hedef maliyet oranı %30)
Yalnız yiyecek-içecek malzeme maliyetidir; işçilik, enerji, kira ve servis dahil değildir. Hedef oran bir varsayımdır; konseptinize göre değiştirin.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


RECETE = {"kalem": ("menu kalemi", "urun", "recete", "yemek", "kalem"), "malzeme": ("malzeme", "hammadde", "bilesen"), "miktar": ("miktar", "net miktar", "gramaj"), "birim": ("birim",), "fire": ("fire", "fire orani")}
FIYAT = {"malzeme": RECETE["malzeme"], "fiyat": ("alis fiyati", "fiyat", "birim fiyat"), "birim": ("birim", "alis birimi")}
MENU = {"kalem": RECETE["kalem"], "fiyat": ("satis fiyati", "fiyat"), "porsiyon": ("porsiyon sayisi", "porsiyon", "recete porsiyonu")}
BIRIM = {"g": ("agirlik", 1), "gr": ("agirlik", 1), "gram": ("agirlik", 1), "kg": ("agirlik", 1000), "ml": ("hacim", 1), "cl": ("hacim", 10), "l": ("hacim", 1000), "lt": ("hacim", 1000), "litre": ("hacim", 1000),
         "adet": ("adet", 1), "ad": ("adet", 1), "tane": ("adet", 1)}
Z = Decimal(0)


def olcu(x) -> Decimal | None:
    """Miktar: 12,5 = 12.5; 1.250,5 = bin iki yüz elli buçuk. Tek başına nokta ondalık ayırıcıdır."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    else:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def oku(recete_yol: Path, fiyat_yol: Path, menu_yol: Path | None = None) -> tuple[list[dict], dict, dict, list[dict]]:
    uy, recete, fiyat, menu = [], [], {}, {}
    for r in kayitlar(recete_yol, RECETE, ("kalem", "malzeme", "miktar", "birim")):
        k, m = metin(r["kalem"]), metin(r["malzeme"])
        if not k and not m:
            continue
        mik, f, b = olcu(r["miktar"]), olcu(r["fire"]) if metin(r["fire"]) else Z, katla(r["birim"])
        if not k or not m or mik is None or mik <= 0 or f is None or not 0 <= f < 100 or b not in BIRIM:
            uy.append(uyari("Yüksek", "Reçete satırı okunamadı", f"{k or '?'} · {m or '?'} (satır {r['_satir']})", f"Miktar '{metin(r['miktar'])}', birim '{metin(r['birim'])}', fire '{metin(r['fire'])}'. "
                            f"Birim şunlardan biri olmalı: {', '.join(BIRIM)}. Satır alınmadı; kalemin maliyeti eksik."))
            continue
        recete.append({"kalem": k, "malzeme": m, "miktar": mik, "birim": b, "fire": f})
    for r in kayitlar(fiyat_yol, FIYAT, ("malzeme", "fiyat", "birim")):
        m = metin(r["malzeme"])
        if not m:
            continue
        f, b = para(r["fiyat"]), katla(r["birim"])
        if f is None or f < 0 or b not in BIRIM:
            uy.append(uyari("Yüksek", "Fiyat okunamadı", m, f"Alış fiyatı '{metin(r['fiyat'])}', birim '{metin(r['birim'])}'. Bu malzeme fiyatsız sayıldı."))
            continue
        if katla(m) in fiyat:
            uy.append(uyari("Orta", "Mükerrer fiyat", m, "Malzeme fiyat listesinde birden fazla kez var. İlk satır kullanıldı."))
            continue
        fiyat[katla(m)] = {"fiyat": f, "birim": b}
    if menu_yol:
        for r in kayitlar(menu_yol, MENU, ("kalem",)):
            k = metin(r["kalem"])
            if not k:
                continue
            f, p = (para(r["fiyat"]) if metin(r["fiyat"]) else None), (olcu(r["porsiyon"]) if metin(r["porsiyon"]) else Decimal(1))
            if (metin(r["fiyat"]) and (f is None or f <= 0)) or p is None or p <= 0:
                uy.append(uyari("Orta", "Menü satırı okunamadı", k, f"Satış fiyatı '{metin(r['fiyat'])}', porsiyon sayısı '{metin(r['porsiyon'])}'. Kalem 1 porsiyon, satış fiyatsız hesaplandı."))
                continue
            menu[katla(k)] = {"fiyat": f, "porsiyon": p}
    return recete, fiyat, menu, uy


def hesapla(recete: list[dict], fiyat: dict, menu: dict | None = None, hedef: Decimal = Decimal(30), kdv: Decimal = Z) -> dict:
    if not recete:
        raise ValueError("Hesaplanacak reçete satırı yok.")
    menu, uy, grup = menu or {}, [], {}
    for s in recete:
        grup.setdefault(katla(s["kalem"]), {"kalem": s["kalem"], "satirlar": []})["satirlar"].append(s)
    kalemler, satirlar = [], []
    for kk, g in grup.items():
        toplam, eksik, ara = Z, [], []
        for s in g["satirlar"]:
            f = fiyat.get(katla(s["malzeme"]))
            brut, m, neden = s["miktar"] / (1 - s["fire"] / 100), None, ""
            if f is None:
                neden = "fiyat yok"
            elif BIRIM[f["birim"]][0] != BIRIM[s["birim"]][0]:
                neden = f"birim uyuşmuyor ({s['birim']} / {f['birim']})"
            else:
                m = brut * BIRIM[s["birim"]][1] / BIRIM[f["birim"]][1] * f["fiyat"]
                toplam += m
            if neden:
                eksik.append(f"{s['malzeme']} ({neden})")
            ara.append({"kalem": g["kalem"], "malzeme": s["malzeme"], "miktar": s["miktar"], "birim": s["birim"], "fire": s["fire"], "brut": yuvarla(brut, "0.001"), "fiyat": f["fiyat"] if f else None,
                        "fiyat_birim": f["birim"] if f else "", "ham": m})
        for a in ara:
            m = a.pop("ham")
            satirlar.append(a | {"maliyet": yuvarla(m) if m is not None else None, "pay": yuvarla(m / toplam * 100) if m is not None and toplam else None})
        if eksik:
            uy.append(uyari("Yüksek", "Maliyeti hesaplanamayan malzeme", g["kalem"], f"{'; '.join(eksik)}. Porsiyon maliyeti bu malzemeler olmadan hesaplandı; gerçek maliyet daha yüksektir."))
        bilgi = menu.get(kk, {"fiyat": None, "porsiyon": Decimal(1)})
        porsiyon = toplam / bilgi["porsiyon"]
        net = bilgi["fiyat"] / (1 + kdv / 100) if bilgi["fiyat"] else None
        oran = porsiyon / net * 100 if net else None
        hedef_fiyat = porsiyon / (hedef / 100) * (1 + kdv / 100)
        durum = "Eksik maliyet" if eksik else "Satış fiyatı yok" if oran is None else "Hedefin üzerinde" if oran > hedef else "Hedefte"
        if durum == "Hedefin üzerinde":
            uy.append(uyari("Orta", "Maliyet oranı hedefin üzerinde", g["kalem"], f"Porsiyon maliyeti {yuvarla(porsiyon):,.2f}, KDV hariç fiyat {yuvarla(net):,.2f}: maliyet oranı %{yuvarla(oran)} (hedef %{hedef.normalize():f}). "
                            f"Hedef için satış fiyatı {yuvarla(hedef_fiyat):,.2f} olmalı veya reçete gözden geçirilmeli."))
        kalemler.append({"kalem": g["kalem"], "malzeme": len(ara), "recete": yuvarla(toplam), "porsiyon_sayisi": bilgi["porsiyon"], "porsiyon": yuvarla(porsiyon), "fiyat": bilgi["fiyat"], "net": yuvarla(net) if net else None,
                         "oran": yuvarla(oran) if oran is not None else None, "kar": yuvarla(net - porsiyon) if net else None, "hedef_fiyat": yuvarla(hedef_fiyat), "durum": durum, "eksik": eksik})
    for kk in menu:
        if kk not in grup:
            uy.append(uyari("Bilgi", "Reçetesi olmayan menü kalemi", kk, "Menü dosyasında var ama reçete dosyasında satırı yok."))
    return {"kalemler": kalemler, "satirlar": satirlar, "ayar": (hedef, kdv), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    hedef, kdv = r["ayar"]
    ws = wb.active
    ws.title = "Porsiyon Maliyeti"
    _baslik(ws, ["Menü Kalemi", "Malzeme Sayısı", "Reçete Maliyeti", "Porsiyon Sayısı", "Porsiyon Maliyeti", "Satış Fiyatı", "KDV Hariç Fiyat", "Maliyet Oranı %", "Brüt Kâr", "Hedef Fiyat (KDV dahil)", "Durum", "Not (siz doldurun)"],
            (30, 9, 14, 9, 14, 12, 13, 10, 12, 14, 18, 30))
    renk = {"Hedefin üzerinde": KIRMIZI, "Eksik maliyet": KIRMIZI, "Hedefte": YESIL, "Satış fiyatı yok": SARI}
    for k in r["kalemler"]:
        _satir(ws, [k["kalem"], k["malzeme"], k["recete"], k["porsiyon_sayisi"], k["porsiyon"], k["fiyat"], k["net"], k["oran"], k["kar"], k["hedef_fiyat"], k["durum"], ""], para_s=(3, 5, 6, 7, 8, 9, 10), dolgu={11: renk[k["durum"]]})
    ws.append([])
    ws.append([f"Hedef maliyet oranı %{hedef.normalize():f}; hedef fiyat = porsiyon maliyeti ÷ hedef oran" + (f", %{kdv.normalize():f} KDV eklenmiş." if kdv else " (fiyatlar KDV hariç kabul edildi).")
               + " Yalnız malzeme maliyetidir; işçilik, enerji ve genel giderler dahil değildir."])
    ws = wb.create_sheet("Malzeme Detayı")
    _baslik(ws, ["Menü Kalemi", "Malzeme", "Net Miktar", "Birim", "Fire %", "Brüt Miktar", "Alış Fiyatı", "Alış Birimi", "Maliyet", "Maliyet Payı %"], (30, 26, 11, 7, 8, 12, 12, 10, 12, 11))
    for s in r["satirlar"]:
        _satir(ws, [s["kalem"], s["malzeme"], s["miktar"], s["birim"], s["fire"] or None, s["brut"], s["fiyat"], s["fiyat_birim"], s["maliyet"], s["pay"]], para_s=(7, 9, 10), dolgu={9: KIRMIZI if s["maliyet"] is None else None})
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(recete_yol: Path, fiyat_yol: Path, cikti: Path, menu_yol: Path | None = None, **ayar) -> dict:
    recete, fiyat, menu, uy = oku(recete_yol, fiyat_yol, menu_yol)
    r = hesapla(recete, fiyat, menu, **ayar)
    uy += r["uyarilar"]
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Standart reçete ve alış fiyatlarından porsiyon maliyetini ve maliyet oranını hesaplar.")
    p.add_argument("--receteler", type=Path, default=ORNEK / "receteler.csv", help="Menü Kalemi, Malzeme, Miktar, Birim, Fire %%")
    p.add_argument("--fiyatlar", type=Path, help="Malzeme, Alış Fiyatı, Birim")
    p.add_argument("--menu", type=Path, help="İsteğe bağlı: Menü Kalemi, Satış Fiyatı, Porsiyon Sayısı")
    p.add_argument("--hedef", default="30", help="Hedef maliyet oranı %% (varsayılan 30)")
    p.add_argument("--kdv", default="0", help="Satış fiyatları KDV dahilse oran, %% (varsayılan 0)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "porsiyon_maliyeti.xlsx")
    a = p.parse_args(argv)
    ornek = a.receteler == ORNEK / "receteler.csv"
    fiyat, menu = a.fiyatlar or (ORNEK / "fiyatlar.csv" if ornek else None), a.menu or (ORNEK / "menu.csv" if ornek else None)
    h, k = para(a.hedef), para(a.kdv)
    if fiyat is None or h is None or not 0 < h <= 100 or k is None or not 0 <= k < 100:
        print("[X] --fiyatlar dosyası verilmeli; --hedef 0–100, --kdv 0 ile 100 arasında olmalı")
        return 2
    for y in (a.receteler, fiyat, menu):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.receteler, fiyat, a.cikti, menu, hedef=h, kdv=k)
    except ValueError as hata:
        print(f"[X] {hata}")
        return 1
    for x in r["kalemler"]:
        ek = f" · maliyet oranı %{x['oran']}" if x["oran"] is not None else ""
        print(f"[{'OK' if x['durum'] == 'Hedefte' else '!'}] {x['kalem']}: porsiyon maliyeti {x['porsiyon']:,.2f}{ek} · {x['durum']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
