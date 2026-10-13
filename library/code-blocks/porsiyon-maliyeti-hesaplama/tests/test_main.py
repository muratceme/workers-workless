import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def s(kalem, malzeme, miktar, birim, fire=0):
    return {"kalem": kalem, "malzeme": malzeme, "miktar": D(miktar), "birim": birim, "fire": D(fire)}


def f(fiyat, birim):
    return {"fiyat": D(fiyat), "birim": birim}


KOFTE = lambda: [s("Köfte", "Kıyma", 150, "g"), s("Köfte", "Soğan", 30, "g", 10), s("Köfte", "Ekmek", 1, "adet"), s("Köfte", "Yağ", 10, "ml")]   # noqa: E731
FIYAT = {"kiyma": f(600, "kg"), "sogan": f(40, "kg"), "ekmek": f(5, "adet"), "yag": f(200, "l")}


class HesapTesti(unittest.TestCase):
    def test_porsiyon_maliyeti(self):
        r = main.hesapla(KOFTE(), FIYAT, {"kofte": {"fiyat": D(350), "porsiyon": D(1)}}, kdv=D(10))
        self.assertEqual([(x["malzeme"], x["brut"], x["maliyet"], x["pay"]) for x in r["satirlar"]],
                         [("Kıyma", D("150.000"), D("90.00"), D("91.53")), ("Soğan", D("33.333"), D("1.33"), D("1.36")), ("Ekmek", D("1.000"), D("5.00"), D("5.08")), ("Yağ", D("10.000"), D("2.00"), D("2.03"))])
        k = r["kalemler"][0]
        self.assertEqual((k["recete"], k["porsiyon"], k["net"], k["oran"], k["kar"], k["hedef_fiyat"], k["durum"], [u["tur"] for u in r["uyarilar"]]),
                         (D("98.33"), D("98.33"), D("318.18"), D("30.90"), D("219.85"), D("360.56"), "Hedefin üzerinde", ["Maliyet oranı hedefin üzerinde"]))
        k = main.hesapla(KOFTE(), FIYAT, {"kofte": {"fiyat": D(350), "porsiyon": D(1)}}, hedef=D(35))["kalemler"][0]
        self.assertEqual((k["net"], k["oran"], k["hedef_fiyat"], k["durum"]), (D("350.00"), D("28.10"), D("280.95"), "Hedefte"))

    def test_birimler_ve_porsiyon_sayisi(self):
        r = main.hesapla([s("Çorba", "Mercimek", 1, "kg"), s("Çorba", "Tereyağı", 100, "gr"), s("Çorba", "Su", "2.5", "lt"), s("Çorba", "Krema", 20, "cl"), s("Çorba", "Yumurta", 2, "tane")],
                         {"mercimek": f(80, "kg"), "tereyagi": f("0.5", "g"), "su": f(4, "l"), "krema": f("0.3", "ml"), "yumurta": f(6, "adet")}, {"corba": {"fiyat": D(130), "porsiyon": D(10)}})
        self.assertEqual(([x["maliyet"] for x in r["satirlar"]], r["kalemler"][0]["recete"], r["kalemler"][0]["porsiyon"], r["kalemler"][0]["oran"], r["kalemler"][0]["durum"]),
                         ([D("80.00"), D("50.00"), D("10.00"), D("60.00"), D("12.00")], D("212.00"), D("21.20"), D("16.31"), "Hedefte"))

    def test_eksikler(self):
        r = main.hesapla([s("A", "Limon", 2, "adet"), s("A", "Un", 500, "g"), s("A", "Süt", 200, "g"), s("A", "Tuz", 5, "g"), s("B", "Un", 1, "kg")], {"limon": f(60, "kg"), "un": f(30, "kg"), "sut": f(40, "l")},
                         {"b": {"fiyat": None, "porsiyon": D(4)}, "c": {"fiyat": D(5), "porsiyon": D(1)}})
        self.assertEqual([(k["kalem"], k["recete"], k["porsiyon"], k["oran"], k["durum"], k["eksik"]) for k in r["kalemler"]],
                         [("A", D("15.00"), D("15.00"), None, "Eksik maliyet", ["Limon (birim uyuşmuyor (adet / kg))", "Süt (birim uyuşmuyor (g / l))", "Tuz (fiyat yok)"]), ("B", D("30.00"), D("7.50"), None, "Satış fiyatı yok", [])])
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]), [("Maliyeti hesaplanamayan malzeme", "A"), ("Reçetesi olmayan menü kalemi", "c")])
        with self.assertRaises(ValueError):
            main.hesapla([], {})


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "p.xlsx"
            r = main.calistir(ORNEK / "receteler.csv", ORNEK / "fiyatlar.csv", cikti, ORNEK / "menu.csv")
            self.assertEqual([(k["kalem"], k["recete"], k["porsiyon"], k["oran"], k["kar"], k["hedef_fiyat"], k["durum"]) for k in r["kalemler"]],
                             [("Izgara Köfte", D("98.33"), D("98.33"), D("30.73"), D("221.67"), D("327.78"), "Hedefin üzerinde"), ("Mercimek Çorbası", D("138.89"), D("13.89"), D("11.57"), D("106.11"), D("46.30"), "Eksik maliyet"),
                              ("Mevsim Salata", D("16.11"), D("16.11"), D("26.84"), D("43.89"), D("53.68"), "Hedefte")])
            self.assertEqual([(u["tur"], u["kim"]) for u in r["uyarilar"]], [("Maliyeti hesaplanamayan malzeme", "Mercimek Çorbası"), ("Maliyet oranı hedefin üzerinde", "Izgara Köfte")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Porsiyon Maliyeti", "Malzeme Detayı", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            ry, fy, my = Path(tmp) / "r.csv", Path(tmp) / "f.csv", Path(tmp) / "m.csv"
            ry.write_text("Yemek;Malzeme;Gramaj;Birim;Fire\nA;Un;12,5;GR;%2\nA;Su;1;bardak;\nA;Tuz;0;g;\nA;;5;g;\n", encoding="utf-8")
            fy.write_text("Malzeme;Fiyat;Birim\nUn;1.250,50;KG\nUn;5;kg\nSu;1;kova\n", encoding="utf-8")
            my.write_text("Yemek;Fiyat;Porsiyon\nA;150;4\nB;x;1\nC;;\n", encoding="utf-8")
            recete, fiyat, menu, uy = main.oku(ry, fy, my)
            self.assertEqual((recete, fiyat, menu, sorted(u["tur"] for u in uy)),
                             ([{"kalem": "A", "malzeme": "Un", "miktar": D("12.5"), "birim": "gr", "fire": D(2)}], {"un": {"fiyat": D("1250.50"), "birim": "kg"}}, {"a": {"fiyat": D(150), "porsiyon": D(4)}, "c": {"fiyat": None, "porsiyon": D(1)}},
                              ["Fiyat okunamadı", "Menü satırı okunamadı", "Mükerrer fiyat"] + ["Reçete satırı okunamadı"] * 3))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--receteler", str(ry), "--fiyatlar", str(fy), "--cikti", c]
            self.assertEqual(main.main(arg + ["--menu", str(my), "--hedef", "35", "--kdv", "10"]), 0)
            self.assertEqual(main.main(arg), 0)
            for ek in (["--hedef", "0"], ["--kdv", "100"]):
                self.assertEqual(main.main(arg + ek), 2)
            self.assertEqual(main.main(["--receteler", str(ry), "--cikti", c]), 2)
            self.assertEqual(main.main(arg + ["--menu", str(Path(tmp) / "yok.csv")]), 1)
            ry.write_text("Yemek;Malzeme\nA;Un\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
