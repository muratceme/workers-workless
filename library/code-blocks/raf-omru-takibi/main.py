"""
Raf Ömrü Takibi — Gerisi Bende kod bloğu
Kalite Kontrol ve Laboratuvar › Kalite Kontrol Teknisyeni
Stok ve sevkiyat verisinden son tüketim tarihine kalan süreyi hesaplar, müşteri kabul kriterini karşılamayan partileri listeler.
  Son tüketim tarihi = STT / TETT sütunu; yoksa üretim tarihi + raf ömrü (gün)
  Kalan ömür         = son tüketim tarihi − değerlendirme günü (gün)  ·  Kalan ömür % = kalan ÷ (son tüketim − üretim) × 100
  Müşteri kriteri    = müşteri dosyasındaki en az kalan ömür % ya da en az kalan gün; ürüne özel kural genel kuraldan önce gelir;
                       müşteri kuralı yoksa --kabul (%66) kullanılır
  Stok durumu        = Süresi dolmuş · Kabul altı (varsayılan kriteri karşılamıyor) · Yaklaşan (kalan ≤ --yakin, 30 gün) · Uygun
  Sevkiyat kontrolü  = her sevk satırı sevk tarihindeki kalan ömre göre müşterinin kriteriyle karşılaştırılır
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

STOK = {"kod": ("urun kodu", "stok kodu", "kod", "malzeme kodu"), "ad": ("urun adi", "urun", "aciklama"), "parti": ("parti no", "lot no", "parti", "lot"),
        "uretim": ("uretim tarihi", "uretim", "imalat tarihi"), "stt": ("son tuketim tarihi", "stt", "tett", "son kullanma tarihi", "skt"), "omur": ("raf omru", "raf omru gun", "omur"),
        "miktar": ("miktar", "stok", "adet", "koli"), "depo": ("depo", "lokasyon")}
MUSTERI = {"musteri": ("musteri", "musteri adi", "cari"), "kod": ("urun kodu", "stok kodu", "kod"), "yuzde": ("kabul kriteri", "en az kalan omur", "kalan omur yuzde", "kabul yuzdesi"),
           "gun": ("en az kalan gun", "kabul gun", "kalan gun")}
SEVK = {"tarih": ("sevk tarihi", "tarih", "planlanan sevk"), "musteri": ("musteri", "musteri adi", "cari"), "kod": ("urun kodu", "stok kodu", "kod"), "parti": ("parti no", "lot no", "parti", "lot"),
        "miktar": ("miktar", "adet", "koli")}
SIFIR = Decimal(0)


def oku_stok(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, STOK, ("kod", "parti")):
        kod, parti = metin(r["kod"]), metin(r["parti"])
        if not kod:
            continue
        ur = tarih(r["uretim"]) if metin(r["uretim"]) else None
        stt = tarih(r["stt"]) if metin(r["stt"]) else None
        omur = para(r["omur"]) if metin(r["omur"]) else None
        miktar = para(r["miktar"]) if metin(r["miktar"]) else None
        hata = (metin(r["uretim"]) and ur is None) or (metin(r["stt"]) and stt is None) or (metin(r["omur"]) and (omur is None or omur <= 0)) or (metin(r["miktar"]) and (miktar is None or miktar < 0))
        if stt is None and ur and omur:
            stt = ur + timedelta(days=int(omur))
        if hata or stt is None:
            uy.append(uyari("Yüksek", "Stok satırı okunamadı", f"{kod} · {parti}", f"Satır {r['_satir']}: son tüketim tarihi yok ya da üretim tarihi / raf ömrü / miktar okunamadı. Parti alınmadı."))
            continue
        if ur and stt <= ur:
            uy.append(uyari("Yüksek", "Son tüketim üretimden önce", f"{kod} · {parti}", f"Üretim {ur:%d.%m.%Y}, son tüketim {stt:%d.%m.%Y}. Parti alınmadı."))
            continue
        l.append({"kod": kod, "ad": metin(r["ad"]), "parti": parti, "uretim": ur, "stt": stt, "miktar": miktar, "depo": metin(r["depo"])})
    return l, uy


def oku_musteri(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, MUSTERI, ("musteri",)):
        m = metin(r["musteri"])
        if not m:
            continue
        y, g = (para(r["yuzde"]) if metin(r["yuzde"]) else None), (para(r["gun"]) if metin(r["gun"]) else None)
        if (y is None and g is None) or (metin(r["yuzde"]) and (y is None or not 0 <= y <= 100)) or (metin(r["gun"]) and (g is None or g < 0)):
            uy.append(uyari("Yüksek", "Müşteri kuralı okunamadı", m, f"Satır {r['_satir']}: en az kalan ömür yüzdesi (0-100) ya da en az kalan gün gerekli."))
            continue
        l.append({"musteri": m, "kod": metin(r["kod"]), "yuzde": y, "gun": g})
    return l, uy


def oku_sevk(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SEVK, ("tarih", "musteri", "parti")):
        m, parti = metin(r["musteri"]), metin(r["parti"])
        if not m and not parti:
            continue
        t = tarih(r["tarih"])
        if t is None or not m or not parti:
            uy.append(uyari("Yüksek", "Sevk satırı okunamadı", f"{m} · {parti}", f"Satır {r['_satir']}: sevk tarihi, müşteri ya da parti okunamadı."))
            continue
        l.append({"tarih": t, "musteri": m, "kod": metin(r["kod"]), "parti": parti, "miktar": para(r["miktar"]) if metin(r["miktar"]) else None})
    return l, uy


def kalan(p: dict, gun: date) -> tuple[int, Decimal | None]:
    k = (p["stt"] - gun).days
    return k, yuvarla(Decimal(k) * 100 / (p["stt"] - p["uretim"]).days) if p["uretim"] else None


def kural_bul(kurallar: list[dict], musteri: str, kod: str) -> dict | None:
    m = [k for k in kurallar if katla(k["musteri"]) == katla(musteri)]
    return next((k for k in m if k["kod"] and katla(k["kod"]) == katla(kod)), None) or next((k for k in m if not k["kod"]), None)


def uygun_mu(p: dict, gun: date, kural: dict | None, kabul: Decimal) -> tuple[bool, str]:
    k, yz = kalan(p, gun)
    if k < 0:
        return False, "süresi dolmuş"
    if kural is None or (kural["yuzde"] is None and kural["gun"] is None):
        if yz is None:
            return True, "üretim tarihi yok, yüzde hesaplanamadı"
        return yz >= kabul, f"kalan %{yz}, kriter %{kabul} (varsayılan)"
    if kural["yuzde"] is not None and (yz is None or yz < kural["yuzde"]):
        return False, f"kalan %{yz}, kriter %{kural['yuzde']}"
    if kural["gun"] is not None and k < kural["gun"]:
        return False, f"kalan {k} gün, kriter {kural['gun']:,.0f} gün"
    return True, (f"kalan %{yz}" if yz is not None else f"kalan {k} gün")


def takip(stok: list[dict], gun: date, kurallar: list[dict] | None = None, sevkler: list[dict] | None = None, kabul: Decimal = Decimal(66), yakin: int = 30) -> dict:
    if not stok:
        raise ValueError("Stok kaydı yok.")
    uy, kurallar, sonuc = [], kurallar or [], []
    musteriler = sorted({k["musteri"] for k in kurallar}, key=katla)
    for p in sorted(stok, key=lambda p: (p["stt"], katla(p["kod"]), p["parti"])):
        k, yz = kalan(p, gun)
        if k < 0:
            durum = "Süresi dolmuş"
            uy.append(uyari("Yüksek", "Süresi dolmuş parti", f"{p['kod']} · {p['parti']}", f"{p['ad'] or p['kod']}: son tüketim {p['stt']:%d.%m.%Y}" + (f", {p['miktar']:,.0f} birim" if p["miktar"] else "") + ". Sevkiyattan ayırın."))
        elif yz is not None and yz < kabul:
            durum = "Kabul altı"
        elif k <= yakin:
            durum = "Yaklaşan"
        else:
            durum = "Uygun"
        uygun = [m for m in musteriler if (kr := kural_bul(kurallar, m, p["kod"])) is not None and uygun_mu(p, gun, kr, kabul)[0]]
        sonuc.append({**p, "kalan": k, "yuzde": yz, "durum": durum, "musteriler": ", ".join(uygun)})
    partiler = {(katla(p["kod"]), katla(p["parti"])): p for p in stok}
    parti_kod = {katla(p["parti"]): p for p in stok}
    sevk_sonuc = []
    for s in sevkler or []:
        p = partiler.get((katla(s["kod"]), katla(s["parti"]))) or (parti_kod.get(katla(s["parti"])) if not s["kod"] else None)
        kim = f"{s['musteri']} · {s['parti']}"
        if p is None:
            uy.append(uyari("Orta", "Sevk edilen parti stokta yok", kim, f"{s['tarih']:%d.%m.%Y} sevkinin partisi stok listesinde bulunamadı."))
            sevk_sonuc.append({**s, "kalan": None, "yuzde": None, "sonuc": "Parti bulunamadı", "neden": ""})
            continue
        kr = kural_bul(kurallar, s["musteri"], p["kod"])
        if kr is None and kurallar:
            uy.append(uyari("Bilgi", "Müşteri kuralı yok", s["musteri"], f"{p['kod']} için kural yok; varsayılan %{kabul} kullanıldı."))
        ok, neden = uygun_mu(p, s["tarih"], kr, kabul)
        k, yz = kalan(p, s["tarih"])
        sevk_sonuc.append({**s, "kod": p["kod"], "kalan": k, "yuzde": yz, "sonuc": "Uygun" if ok else "Uygun değil", "neden": neden})
        if not ok:
            uy.append(uyari("Yüksek", "Sevkiyat kabul kriterini karşılamıyor", kim, f"{s['tarih']:%d.%m.%Y} · {p['kod']}: {neden}."))
    ozet = {d: sum(1 for s in sonuc if s["durum"] == d) for d in ("Süresi dolmuş", "Kabul altı", "Yaklaşan", "Uygun")}
    return {"stok": sonuc, "sevk": sevk_sonuc, "ozet": ozet, "uyarilar": uy}


DURUM_RENK = {"Süresi dolmuş": KIRMIZI, "Kabul altı": SARI, "Yaklaşan": MAVI, "Uygun": YESIL}


def rapor_yaz(cikti: Path, r: dict, gun: date, kabul: Decimal, yakin: int) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Partiler"
    _baslik(ws, ["Durum", "Ürün Kodu", "Ürün Adı", "Parti No", "Üretim", "Son Tüketim", "Kalan Gün", "Kalan Ömür %", "Miktar", "Depo", "Kriterini Karşıladığı Müşteriler"],
            (13, 12, 26, 12, 11, 11, 8, 10, 9, 12, 30), "C2")
    for p in r["stok"]:
        _satir(ws, [p["durum"], p["kod"], p["ad"], p["parti"], p["uretim"], p["stt"], p["kalan"], p["yuzde"], p["miktar"], p["depo"], p["musteriler"]], para_s=(8,), tarih_s=(5, 6),
               dolgu={1: DURUM_RENK[p["durum"]]}, sar=(11,))
    ws.append([])
    ws.append([f"Değerlendirme günü {gun:%d.%m.%Y}. Kabul altı: kalan ömür %{kabul} altında (müşteri kuralı yoksa). Yaklaşan: {yakin} gün ve altı."])
    if r["sevk"]:
        ws = wb.create_sheet("Sevkiyat Kontrolü")
        _baslik(ws, ["Sevk Tarihi", "Müşteri", "Ürün Kodu", "Parti No", "Miktar", "Sevkte Kalan Gün", "Sevkte Kalan Ömür %", "Sonuç", "Açıklama"], (11, 22, 12, 12, 9, 10, 11, 14, 34))
        for s in r["sevk"]:
            _satir(ws, [s["tarih"], s["musteri"], s["kod"], s["parti"], s["miktar"], s["kalan"], s["yuzde"], s["sonuc"], s["neden"]], para_s=(7,), tarih_s=(1,),
                   dolgu={8: YESIL if s["sonuc"] == "Uygun" else KIRMIZI})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(stok_yol: Path, cikti: Path, gun: date, musteri_yol: Path | None = None, sevk_yol: Path | None = None, kabul: Decimal = Decimal(66), yakin: int = 30) -> dict:
    stok, uy = oku_stok(stok_yol)
    kurallar, uy2 = oku_musteri(musteri_yol) if musteri_yol else ([], [])
    sevkler, uy3 = oku_sevk(sevk_yol) if sevk_yol else ([], [])
    r = takip(stok, gun, kurallar, sevkler, kabul, yakin)
    r["uyarilar"] = sirala(uy + uy2 + uy3 + r["uyarilar"])
    rapor_yaz(cikti, r, gun, kabul, yakin)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Son tüketim tarihine kalan süreyi hesaplar, müşteri kabul kriterini karşılamayan partileri listeler.")
    p.add_argument("--stok", type=Path, help="Ürün Kodu, Ürün Adı, Parti No, Üretim Tarihi, Son Tüketim Tarihi (ya da Raf Ömrü), Miktar, Depo")
    p.add_argument("--musteriler", type=Path, help="İsteğe bağlı: Müşteri, Ürün Kodu (boşsa tüm ürünler), Kabul Kriteri %%, En Az Kalan Gün")
    p.add_argument("--sevkiyat", type=Path, help="İsteğe bağlı: Sevk Tarihi, Müşteri, Ürün Kodu, Parti No, Miktar")
    p.add_argument("--tarih", help="Değerlendirme günü GG.AA.YYYY (varsayılan bugün; örnek veride 15.10.2026)")
    p.add_argument("--kabul", default="66", help="Müşteri kuralı yoksa en az kalan ömür yüzdesi (varsayılan 66)")
    p.add_argument("--yakin", type=int, default=30, help="Kalan gün bu değer ve altındaysa yaklaşan sayılır (varsayılan 30)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "raf_omru_takibi.xlsx")
    a = p.parse_args(argv)
    ornek = a.stok is None
    sy = a.stok or ORNEK / "parti_stoklari.csv"
    my = a.musteriler or (ORNEK / "musteri_kabul_kriterleri.csv" if ornek else None)
    vy = a.sevkiyat or (ORNEK / "planlanan_sevkiyat.csv" if ornek else None)
    gun = tarih(a.tarih) if a.tarih else (date(2026, 10, 15) if ornek else date.today())
    kabul = para(a.kabul)
    if gun is None or kabul is None or not 0 <= kabul <= 100 or a.yakin < 0:
        print("[X] --tarih GG.AA.YYYY, --kabul 0-100 arası, --yakin sıfır ya da üstü olmalı")
        return 2
    for y in (sy, my, vy):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(sy, a.cikti, gun, my, vy, kabul, a.yakin)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {len(r['stok'])} parti · süresi dolmuş {o['Süresi dolmuş']} · kabul altı {o['Kabul altı']} · yaklaşan {o['Yaklaşan']} · uygun {o['Uygun']}"
          + (f" · sevk kontrolü {len(r['sevk'])} satır, uygun değil {sum(1 for s in r['sevk'] if s['sonuc'] == 'Uygun değil')}" if r["sevk"] else ""))
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
