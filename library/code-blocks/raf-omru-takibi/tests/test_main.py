import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
GUN = date(2026, 10, 15)


def p(kod, parti, uretim, stt):
    return {"kod": kod, "ad": "", "parti": parti, "uretim": date(2026, *uretim) if uretim else None, "stt": date(2026, *stt), "miktar": D(10), "depo": ""}


def k(m, kod="", yuzde=None, gun=None):
    return {"musteri": m, "kod": kod, "yuzde": None if yuzde is None else D(yuzde), "gun": None if gun is None else D(gun)}


def s(m, parti, kod="U1"):
    return {"tarih": date(2026, 10, 16), "musteri": m, "kod": kod, "parti": parti, "miktar": None}


STOK = [p("U1", "P1", (9, 1), (11, 30)), p("U1", "P2", (10, 1), (12, 30)), p("U2", "P3", (6, 1), (9, 29))]
KURAL = [k("M1", yuzde=66), k("M2", "U1", gun=30), k("M3", "U2", yuzde=50)]


class TakipTesti(unittest.TestCase):
    def test_stok_ve_sevk(self):
        r = main.takip(STOK, GUN, KURAL, [s("M1", "P1"), s("M2", "P1"), s("M1", "P2"), s("M9", "P2"), s("M1", "P9")])
        self.assertEqual([(x["parti"], x["kalan"], x["yuzde"], x["durum"], x["musteriler"]) for x in r["stok"]],
                         [("P3", -16, D("-13.33"), "Süresi dolmuş", ""), ("P1", 46, D("51.11"), "Kabul altı", "M2"), ("P2", 76, D("84.44"), "Uygun", "M1, M2")])
        self.assertEqual([(x["musteri"], x["parti"], x["sonuc"], x["neden"]) for x in r["sevk"]],
                         [("M1", "P1", "Uygun değil", "kalan %50.00, kriter %66"), ("M2", "P1", "Uygun", "kalan %50.00"), ("M1", "P2", "Uygun", "kalan %83.33"),
                          ("M9", "P2", "Uygun", "kalan %83.33, kriter %66 (varsayılan)"), ("M1", "P9", "Parti bulunamadı", "")])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Bilgi", "Müşteri kuralı yok", "M9"), ("Orta", "Sevk edilen parti stokta yok", "M1 · P9"), ("Yüksek", "Sevkiyat kabul kriterini karşılamıyor", "M1 · P1"),
                          ("Yüksek", "Süresi dolmuş parti", "U2 · P3")])

    def test_kurallar(self):
        r = main.takip(STOK, GUN, kabul=D(50), yakin=50)
        self.assertEqual([x["durum"] for x in r["stok"]], ["Süresi dolmuş", "Yaklaşan", "Uygun"])
        self.assertEqual(main.kural_bul([k("A", yuzde=60), k("a", "X", gun=5)], "A", "x")["gun"], D(5))
        self.assertEqual(main.uygun_mu(p("U", "P", None, (12, 1)), GUN, None, D(66)), (True, "üretim tarihi yok, yüzde hesaplanamadı"))
        self.assertEqual(main.uygun_mu(p("U", "P", None, (12, 1)), GUN, k("A", yuzde=10), D(66)), (False, "kalan %None, kriter %10"))
        with self.assertRaises(ValueError):
            main.takip([], GUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "parti_stoklari.csv", cikti, GUN, ORNEK / "musteri_kabul_kriterleri.csv", ORNEK / "planlanan_sevkiyat.csv")
            self.assertEqual((r["ozet"], [x["sonuc"] for x in r["sevk"]]), ({"Süresi dolmuş": 1, "Kabul altı": 3, "Yaklaşan": 2, "Uygun": 3},
                             ["Uygun", "Uygun değil", "Uygun", "Uygun değil", "Uygun değil", "Uygun değil", "Uygun", "Uygun"]))
            self.assertEqual([(u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Sevkiyat kabul kriterini karşılamıyor", "Toptancı C · B260320"), ("Sevkiyat kabul kriterini karşılamıyor", "Zincir Market A · Y261005"),
                              ("Sevkiyat kabul kriterini karşılamıyor", "Zincir Market B · P260615"), ("Sevkiyat kabul kriterini karşılamıyor", "Zincir Market B · P260901"),
                              ("Süresi dolmuş parti", "YG-500 · Y260910"), ("Müşteri kuralı yok", "Bakkal E")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Partiler", "Sevkiyat Kontrolü", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, m, v = Path(tmp) / "s.csv", Path(tmp) / "m.csv", Path(tmp) / "v.csv"
            y.write_text("Stok Kodu;Lot No;Üretim Tarihi;Raf Ömrü;SKT;Koli\nA;L1;01.10.2026;90;;1.200\nA;L2;x;;;\nA;L3;;;;\nA;L4;10.10.2026;;01.10.2026;\n;L5;;;01.01.2027;\n", encoding="utf-8")
            m.write_text("Müşteri;Kabul Kriteri\nX;70\nY;\nZ;150\n", encoding="utf-8")
            v.write_text("Sevk Tarihi;Müşteri;Lot No\n20.10.2026;X;L1\nx;X;L1\n", encoding="utf-8")
            l, uy = main.oku_stok(y)
            self.assertEqual(([(x["parti"], x["stt"], x["miktar"]) for x in l], len(uy)), ([("L1", date(2026, 12, 30), D(1200))], 3))
            self.assertEqual((len(main.oku_musteri(m)[0]), len(main.oku_musteri(m)[1]), len(main.oku_sevk(v)[0]), len(main.oku_sevk(v)[1])), (1, 2, 1, 1))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--stok", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--musteriler", str(m), "--sevkiyat", str(v), "--tarih", "15.10.2026", "--kabul", "50", "--yakin", "20"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + x) for x in (["--tarih", "x"], ["--kabul", "101"], ["--yakin", "-1"])], [2] * 3)
            self.assertEqual(main.main(arg + ["--sevkiyat", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Stok Kodu;Koli\nA;1\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
