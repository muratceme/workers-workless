"""
Mali Tablo Rasyo Analizi — Workers / Workless kod bloğu
Finans › Finansal Analist

Muhasebe programlarından alınan mizanlardan (Tekdüzen Hesap Planı) özet bilanço ve gelir tablosu oluşturur;
likidite, finansal yapı, faaliyet (devir hızları ve süreleri, nakit dönüşüm süresi) ve kârlılık rasyolarını
dönemler arası karşılaştırmalı hesaplar. Devir hızlarında önceki dönem varsa ortalama bakiye kullanılır.
İnternete bağlanmaz.

Mizan kapanış kayıtlarından ÖNCE alınmalıdır (6'lı hesaplar açık); kapanmış mizanda gelir tablosu boş kalır.
Ana hesap (3 hane) satırı varsa o, yoksa en alt kırılımların toplamı kullanılır; ara seviyeler çift sayılmaz.

Kullanım:
    python main.py                                              # örnek 2025 ve 2026 mizanlarıyla dener
    python main.py --mizan 2025=mizan_2025.xlsx 2026=mizan_2026.xlsx
    python main.py --mizan 2026/09=mizan_eylul.xlsx --gun 273   # dönem 273 gün (9 ay)
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import OrderedDict, defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
SIFIR = Decimal(0)


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def para(x) -> Decimal:
    if x in (None, ""):
        return SIFIR
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("₺", "").replace(" ", "")
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        d = Decimal(s)
    except InvalidOperation:
        return SIFIR
    return -d if neg else d


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                metin = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        # Üstte firma/dönem başlığı olabilir: ayırıcı, ilk satırlarda en çok geçen karakterdir
        ilk = "\n".join(metin.splitlines()[:10])
        ayirici = max(";\t,", key=ilk.count)
        satirlar = list(csv.reader(metin.splitlines(), delimiter=ayirici))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


# ----------------------------------------------------------------------------
# Mizan okuma: ana hesap bakiyeleri (borç − alacak)
# ----------------------------------------------------------------------------

def mizan_oku(yol: Path) -> tuple[dict[str, Decimal], list[str]]:
    s = tablo_oku(yol)
    # başlık satırı: "hesap kodu" içeren ilk satır (üstte firma bilgisi olabilir)
    bi = next((i for i, r in enumerate(s[:15]) if any(kucuk(c) in ("hesap kodu", "hesap no", "kod") for c in r)), 0)
    b = [kucuk(x) for x in s[bi]]
    bul = lambda *adlar: next((i for i, x in enumerate(b) if x in adlar), None)  # noqa: E731
    i_k = bul("hesap kodu", "hesap no", "kod")
    i_bb, i_ab = bul("borç bakiye", "borç bakiyesi", "bakiye borç"), bul("alacak bakiye", "alacak bakiyesi", "bakiye alacak")
    i_b, i_a = bul("borç", "borç toplamı", "toplam borç"), bul("alacak", "alacak toplamı", "toplam alacak")
    i_bak = bul("bakiye", "net bakiye")
    if i_k is None or not ((i_bb is not None and i_ab is not None) or (i_b is not None and i_a is not None) or i_bak is not None):
        raise SystemExit(f"{yol.name}: Hesap Kodu ve Borç/Alacak Bakiye (veya Borç/Alacak ya da Bakiye) sütunları gerekli. "
                         f"Başlıklar: {s[bi]}")
    satirlar: dict[str, Decimal] = {}
    for r in s[bi + 1:]:
        kod = re.sub(r"\s", "", str(r[i_k] or ""))
        if not re.match(r"^\d{3}", kod):
            continue
        al = lambda i: r[i] if i is not None and i < len(r) else None  # noqa: E731
        if i_bb is not None and i_ab is not None:
            net = para(al(i_bb)) - para(al(i_ab))
        elif i_b is not None and i_a is not None:
            net = para(al(i_b)) - para(al(i_a))
        else:
            net = para(al(i_bak))
        satirlar[kod] = satirlar.get(kod, SIFIR) + net
    # Ana hesap varsa o; yoksa yaprak kırılımların toplamı
    ana: dict[str, Decimal] = defaultdict(Decimal)
    kodlar = sorted(satirlar)
    for kod in kodlar:
        ana_kod = kod[:3]
        if kod == ana_kod:
            ana[ana_kod] = satirlar[kod]
        elif ana_kod not in satirlar:
            yaprak = not any(k != kod and k.startswith(kod) and not k[len(kod)].isdigit() for k in kodlar)
            if yaprak:
                ana[ana_kod] += satirlar[kod]
    uyarilar = []
    fark = sum(ana.values(), SIFIR)
    if abs(fark) > 1:
        uyarilar.append(f"{yol.name}: mizan denk değil (borç − alacak = {fark:,.2f}); eksik hesap olabilir")
    return dict(ana), uyarilar


# ----------------------------------------------------------------------------
# Tekdüzen gruplama
# ----------------------------------------------------------------------------

def grup(m: dict[str, Decimal], *onekler: str, alacak: bool = False) -> Decimal:
    t = sum((v for k, v in m.items() if k.startswith(onekler)), SIFIR)
    return -t if alacak else t


def tablolar(m: dict[str, Decimal]) -> tuple[dict, list[str]]:
    uyarilar = []
    g = lambda *o: grup(m, *o, alacak=True)  # noqa: E731  gelir tablosu: alacak − borç (gelir +, gider −)
    gt = {
        "Brüt satışlar": g("60"),
        "Satış indirimleri (−)": g("61"),
    }
    gt["Net satışlar"] = gt["Brüt satışlar"] + gt["Satış indirimleri (−)"]
    gt["Satışların maliyeti (−)"] = g("62")
    gt["Brüt satış kârı"] = gt["Net satışlar"] + gt["Satışların maliyeti (−)"]
    gt["Faaliyet giderleri (−)"] = g("63")
    gt["Esas faaliyet kârı"] = gt["Brüt satış kârı"] + gt["Faaliyet giderleri (−)"]
    gt["Diğer faaliyet gelirleri"] = g("64")
    gt["Diğer faaliyet giderleri (−)"] = g("65")
    gt["Finansman giderleri (−)"] = g("66")
    gt["Olağan kâr"] = gt["Esas faaliyet kârı"] + gt["Diğer faaliyet gelirleri"] + gt["Diğer faaliyet giderleri (−)"] + gt["Finansman giderleri (−)"]
    gt["Olağandışı gelir ve kârlar"] = g("67")
    gt["Olağandışı gider ve zararlar (−)"] = g("68")
    gt["Dönem kârı (vergi öncesi)"] = gt["Olağan kâr"] + gt["Olağandışı gelir ve kârlar"] + gt["Olağandışı gider ve zararlar (−)"]
    gt["Vergi karşılığı (−)"] = g("691")
    gt["Dönem net kârı"] = gt["Dönem kârı (vergi öncesi)"] + gt["Vergi karşılığı (−)"]

    acik_6 = any(v for k, v in m.items() if k.startswith("6") and not k.startswith("69"))
    if not acik_6:
        uyarilar.append("6'lı gelir tablosu hesapları kapalı (mizan kapanış sonrası); gelir tablosu ve kârlılık/faaliyet rasyoları hesaplanamaz")
        donem_kari = grup(m, "59", alacak=True)
    else:
        donem_kari = gt["Dönem net kârı"]
        if grup(m, "590", "591"):
            uyarilar.append("Hem 6'lı hesaplar açık hem 590/591'de bakiye var: dönem kârı çift sayılabilir")
    yedi = grup(m, "7")
    if abs(yedi) > 1:
        uyarilar.append(f"7'li maliyet hesaplarında {yedi:,.2f} bakiye var (yansıtma kayıtları yapılmamış olabilir); bilanço denkliği bozulur")

    bl = OrderedDict()
    bl["Hazır değerler"] = grup(m, "10")
    bl["Menkul kıymetler"] = grup(m, "11")
    bl["Ticari alacaklar"] = grup(m, "12")
    bl["Diğer alacaklar"] = grup(m, "13")
    bl["Stoklar"] = grup(m, "15")
    bl["Yıllara yaygın inşaat ve onarım maliyetleri"] = grup(m, "17")
    bl["Gelecek aylara ait giderler ve gelir tahakkukları"] = grup(m, "18")
    bl["Diğer dönen varlıklar"] = grup(m, "19")
    bl["DÖNEN VARLIKLAR"] = grup(m, "1")
    bl["Ticari alacaklar (uzun vadeli)"] = grup(m, "22")
    bl["Mali duran varlıklar"] = grup(m, "24")
    bl["Maddi duran varlıklar"] = grup(m, "25")
    bl["Maddi olmayan duran varlıklar"] = grup(m, "26")
    bl["Diğer duran varlıklar"] = grup(m, "23", "27", "28", "29")
    bl["DURAN VARLIKLAR"] = grup(m, "2")
    bl["AKTİF TOPLAMI"] = bl["DÖNEN VARLIKLAR"] + bl["DURAN VARLIKLAR"]
    bl["Mali borçlar (kısa vadeli)"] = grup(m, "30", alacak=True)
    bl["Ticari borçlar (kısa vadeli)"] = grup(m, "32", alacak=True)
    bl["Diğer kısa vadeli yabancı kaynaklar"] = grup(m, "33", "34", "35", "36", "37", "38", "39", alacak=True)
    bl["KISA VADELİ YABANCI KAYNAKLAR"] = grup(m, "3", alacak=True)
    bl["Mali borçlar (uzun vadeli)"] = grup(m, "40", alacak=True)
    bl["Diğer uzun vadeli yabancı kaynaklar"] = grup(m, "42", "43", "44", "47", "48", "49", alacak=True)
    bl["UZUN VADELİ YABANCI KAYNAKLAR"] = grup(m, "4", alacak=True)
    bl["Özkaynaklar (dönem kârı hariç)"] = grup(m, "50", "52", "54", "57", "58", alacak=True)
    bl["Dönem net kârı / zararı"] = donem_kari
    bl["ÖZKAYNAKLAR"] = bl["Özkaynaklar (dönem kârı hariç)"] + donem_kari
    bl["PASİF TOPLAMI"] = bl["KISA VADELİ YABANCI KAYNAKLAR"] + bl["UZUN VADELİ YABANCI KAYNAKLAR"] + bl["ÖZKAYNAKLAR"]
    if abs(bl["AKTİF TOPLAMI"] - bl["PASİF TOPLAMI"]) > 1:
        uyarilar.append(f"Bilanço denk değil: aktif {bl['AKTİF TOPLAMI']:,.2f} ≠ pasif {bl['PASİF TOPLAMI']:,.2f}")
    return {"bilanco": bl, "gelir": gt, "acik_6": acik_6}, uyarilar


# ----------------------------------------------------------------------------
# Rasyolar
# ----------------------------------------------------------------------------

def bol(a: Decimal, b: Decimal) -> float | None:
    return None if not b else float(a / b)


# (grup, ad, biçim, formül açıklaması, yorum yönü: +1 yüksek iyi, −1 düşük iyi, 0 nötr, genel referans)
RASYOLAR = [
    ("Likidite", "Cari oran", "x", "Dönen varlıklar / KVYK", 1, "Genelde 1,5 – 2 civarı yeterli kabul edilir"),
    ("Likidite", "Asit-test oranı", "x", "(Dönen varlıklar − Stoklar) / KVYK", 1, "Genelde ≥ 1"),
    ("Likidite", "Nakit oranı", "x", "(Hazır değerler + Menkul kıymetler) / KVYK", 1, "Genelde ≥ 0,2"),
    ("Finansal yapı", "Kaldıraç oranı", "%", "(KVYK + UVYK) / Aktif toplamı", -1, "Genelde ≤ %50 – 60"),
    ("Finansal yapı", "Özkaynak / Aktif", "%", "Özkaynaklar / Aktif toplamı", 1, ""),
    ("Finansal yapı", "Borç / Özkaynak", "x", "(KVYK + UVYK) / Özkaynaklar", -1, ""),
    ("Finansal yapı", "KVYK / Toplam yabancı kaynak", "%", "KVYK / (KVYK + UVYK)", -1, ""),
    ("Finansal yapı", "Duran varlıklar / Özkaynak", "x", "Duran varlıklar / Özkaynaklar", -1, "Genelde ≤ 1"),
    ("Faaliyet", "Ticari alacak devir hızı", "x", "Net satışlar / Ortalama ticari alacaklar", 1, ""),
    ("Faaliyet", "Ortalama tahsil süresi", "gün", "Dönem günü / Alacak devir hızı", -1, ""),
    ("Faaliyet", "Stok devir hızı", "x", "Satışların maliyeti / Ortalama stoklar", 1, ""),
    ("Faaliyet", "Stokta kalma süresi", "gün", "Dönem günü / Stok devir hızı", -1, ""),
    ("Faaliyet", "Ticari borç devir hızı", "x", "Satışların maliyeti / Ortalama ticari borçlar", 0, ""),
    ("Faaliyet", "Ortalama ödeme süresi", "gün", "Dönem günü / Ticari borç devir hızı", 1, ""),
    ("Faaliyet", "Nakit dönüşüm süresi", "gün", "Tahsil süresi + Stokta kalma süresi − Ödeme süresi", -1, ""),
    ("Faaliyet", "Aktif devir hızı", "x", "Net satışlar / Ortalama aktif toplamı", 1, ""),
    ("Kârlılık", "Brüt kâr marjı", "%", "Brüt satış kârı / Net satışlar", 1, ""),
    ("Kârlılık", "Esas faaliyet kâr marjı", "%", "Esas faaliyet kârı / Net satışlar", 1, ""),
    ("Kârlılık", "Net kâr marjı", "%", "Dönem net kârı / Net satışlar", 1, ""),
    ("Kârlılık", "Aktif kârlılığı (ROA)", "%", "Dönem net kârı / Ortalama aktif toplamı", 1, ""),
    ("Kârlılık", "Özkaynak kârlılığı (ROE)", "%", "Dönem net kârı / Ortalama özkaynaklar", 1, ""),
    ("Kârlılık", "Faiz karşılama oranı", "x", "(Vergi öncesi kâr + Finansman giderleri) / Finansman giderleri", 1, "Genelde ≥ 3"),
]


def rasyolar(t: dict, onceki: dict | None, gun: Decimal) -> dict[str, float | None]:
    bl, gt = t["bilanco"], t["gelir"]
    ob = onceki["bilanco"] if onceki else None
    ort = lambda k: (bl[k] + ob[k]) / 2 if ob else bl[k]  # noqa: E731
    kvyk, uvyk = bl["KISA VADELİ YABANCI KAYNAKLAR"], bl["UZUN VADELİ YABANCI KAYNAKLAR"]
    ns, smm = gt["Net satışlar"], -gt["Satışların maliyeti (−)"]
    r = {
        "Cari oran": bol(bl["DÖNEN VARLIKLAR"], kvyk),
        "Asit-test oranı": bol(bl["DÖNEN VARLIKLAR"] - bl["Stoklar"], kvyk),
        "Nakit oranı": bol(bl["Hazır değerler"] + bl["Menkul kıymetler"], kvyk),
        "Kaldıraç oranı": bol(kvyk + uvyk, bl["AKTİF TOPLAMI"]),
        "Özkaynak / Aktif": bol(bl["ÖZKAYNAKLAR"], bl["AKTİF TOPLAMI"]),
        "Borç / Özkaynak": bol(kvyk + uvyk, bl["ÖZKAYNAKLAR"]),
        "KVYK / Toplam yabancı kaynak": bol(kvyk, kvyk + uvyk),
        "Duran varlıklar / Özkaynak": bol(bl["DURAN VARLIKLAR"], bl["ÖZKAYNAKLAR"]),
    }
    if t["acik_6"]:
        adh, sdh, bdh = bol(ns, ort("Ticari alacaklar")), bol(smm, ort("Stoklar")), bol(smm, ort("Ticari borçlar (kısa vadeli)"))
        g = float(gun)
        tahsil = None if not adh else g / adh
        stok = None if not sdh else g / sdh
        odeme = None if not bdh else g / bdh
        vo, fin = gt["Dönem kârı (vergi öncesi)"], -gt["Finansman giderleri (−)"]
        r.update({
            "Ticari alacak devir hızı": adh, "Ortalama tahsil süresi": tahsil,
            "Stok devir hızı": sdh, "Stokta kalma süresi": stok,
            "Ticari borç devir hızı": bdh, "Ortalama ödeme süresi": odeme,
            "Nakit dönüşüm süresi": None if None in (tahsil, stok, odeme) else tahsil + stok - odeme,
            "Aktif devir hızı": bol(ns, ort("AKTİF TOPLAMI")),
            "Brüt kâr marjı": bol(gt["Brüt satış kârı"], ns),
            "Esas faaliyet kâr marjı": bol(gt["Esas faaliyet kârı"], ns),
            "Net kâr marjı": bol(gt["Dönem net kârı"], ns),
            "Aktif kârlılığı (ROA)": bol(gt["Dönem net kârı"], ort("AKTİF TOPLAMI")),
            "Özkaynak kârlılığı (ROE)": bol(gt["Dönem net kârı"], ort("ÖZKAYNAKLAR")),
            "Faiz karşılama oranı": bol(vo + fin, fin),
        })
    return r


def calistir(mizanlar: "OrderedDict[str, Path]", cikti: Path, gunler: dict[str, Decimal] | None = None) -> dict:
    gunler = gunler or {}
    donemler: "OrderedDict[str, dict]" = OrderedDict()
    uyarilar = []
    onceki = None
    for ad, yol in mizanlar.items():
        m, u = mizan_oku(yol)
        t, u2 = tablolar(m)
        uyarilar += u + [f"{ad}: {x}" for x in u2]
        gun = gunler.get(ad, Decimal(365))
        t["rasyolar"] = rasyolar(t, onceki, gun)
        t["ortalama"] = onceki is not None
        donemler[ad] = t
        onceki = t
    _rapor(donemler, cikti, uyarilar)
    return {"donemler": donemler, "uyarilar": uyarilar}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
GRUP_DOLGU = PatternFill("solid", fgColor="EDEDED")
YESIL = PatternFill("solid", fgColor="E3F5E1")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
PARA = "#,##0"
BICIM = {"x": "0.00", "%": "0.0%", "gün": "0"}


def _rapor(donemler, cikti, uyarilar):
    adlar = list(donemler)
    wb = Workbook()
    r = wb.active
    r.title = "Rasyolar"
    r.append(["Grup", "Rasyo", "Formül"] + adlar + (["Değişim", "Yön"] if len(adlar) > 1 else []) + ["Genel referans"])
    for h in r[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    son_grup = None
    for grup_adi, ad, bicim, formul, yon, ref in RASYOLAR:
        if grup_adi != son_grup:
            r.append([grup_adi.upper()])
            for c in range(1, 4 + len(adlar) + 2):
                r.cell(r.max_row, c).fill = GRUP_DOLGU
            r.cell(r.max_row, 1).font = Font(bold=True)
            son_grup = grup_adi
        degerler = [donemler[d]["rasyolar"].get(ad) for d in adlar]
        satir = ["", ad, formul] + degerler
        if len(adlar) > 1:
            a, b = degerler[-2], degerler[-1]
            degisim = None if a is None or b is None else b - a
            iyi = "" if degisim is None or yon == 0 or abs(degisim) < 1e-9 else ("İyileşme" if degisim * yon > 0 else "Kötüleşme")
            satir += [degisim, iyi]
        satir.append(ref)
        r.append(satir)
        for j in range(4, 4 + len(adlar) + (1 if len(adlar) > 1 else 0)):
            r.cell(r.max_row, j).number_format = BICIM[bicim]
        if len(adlar) > 1 and satir[-2]:
            r.cell(r.max_row, 4 + len(adlar) + 1).fill = YESIL if satir[-2] == "İyileşme" else KIRMIZI
    for j, w in enumerate([14, 30, 56] + [13] * len(adlar) + ([12, 12] if len(adlar) > 1 else []) + [34], 1):
        r.column_dimensions[get_column_letter(j)].width = w
    r.freeze_panes = "D2"
    if any(donemler[d]["ortalama"] for d in adlar):
        r.append([])
        r.append(["Not", "Devir hızları ve ROA/ROE'de önceki dönem varsa ortalama bakiye ((önceki + son) / 2), yoksa dönem sonu bakiyesi kullanılır"])
    for u in uyarilar:
        r.append(["Uyarı", u])

    for ad, anahtar in (("Bilanço", "bilanco"), ("Gelir Tablosu", "gelir")):
        ws = wb.create_sheet(ad)
        ws.append(["Kalem"] + adlar + (["Değişim %"] if len(adlar) > 1 else []))
        for h in ws[1]:
            h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        for kalem in donemler[adlar[0]][anahtar]:
            vals = [donemler[d][anahtar][kalem] for d in adlar]
            satir = [kalem] + [float(v) for v in vals]
            if len(adlar) > 1:
                satir.append(None if not vals[-2] else float((vals[-1] - vals[-2]) / abs(vals[-2])))
            ws.append(satir)
            for j in range(2, 2 + len(adlar)):
                ws.cell(ws.max_row, j).number_format = PARA
            if len(adlar) > 1:
                ws.cell(ws.max_row, 2 + len(adlar)).number_format = "0.0%"
            if kalem.isupper() or kalem in ("Net satışlar", "Brüt satış kârı", "Esas faaliyet kârı", "Dönem net kârı"):
                ws.cell(ws.max_row, 1).font = Font(bold=True)
        ws.column_dimensions["A"].width = 48
        for j in range(2, 3 + len(adlar)):
            ws.column_dimensions[get_column_letter(j)].width = 18
        ws.freeze_panes = "B2"

    # Dikey analiz (aktif yüzdeleri) ve grafik
    v = wb.create_sheet("Dikey Analiz")
    v.append(["Kalem"] + adlar)
    for h in v[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for kalem in donemler[adlar[0]]["bilanco"]:
        v.append([kalem] + [bol(donemler[d]["bilanco"][kalem], donemler[d]["bilanco"]["AKTİF TOPLAMI"]) for d in adlar])
        for j in range(2, 2 + len(adlar)):
            v.cell(v.max_row, j).number_format = "0.0%"
    v.column_dimensions["A"].width = 48
    if donemler[adlar[-1]]["acik_6"]:
        g = wb["Gelir Tablosu"]
        ch = BarChart()
        ch.title, ch.height, ch.width = "Net satış, brüt kâr ve net kâr", 8, 18
        satirlar = [i for i, row in enumerate(g.iter_rows(min_col=1, max_col=1, values_only=True), 1)
                    if row[0] in ("Net satışlar", "Brüt satış kârı", "Dönem net kârı")]
        for i in satirlar:
            ch.add_data(Reference(g, min_col=1, max_col=1 + len(adlar), min_row=i, max_row=i), from_rows=True, titles_from_data=True)
        ch.set_categories(Reference(g, min_col=2, max_col=1 + len(adlar), min_row=1, max_row=1))
        g.add_chart(ch, f"{get_column_letter(len(adlar) + 4)}2")
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description="Mizanlardan bilanço, gelir tablosu ve karşılaştırmalı rasyo analizi üretir.")
    ap.add_argument("--mizan", nargs="+", metavar="DÖNEM=DOSYA",
                    help="Dönem adı ve mizan dosyası (.xlsx/.csv), eskiden yeniye: 2025=mizan_2025.xlsx 2026=mizan_2026.xlsx")
    ap.add_argument("--gun", nargs="*", default=[], metavar="GÜN",
                    help="Dönem gün sayıları (devir süreleri için), sırayla; verilmezse 365. Ör. 365 273")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "rasyo_analizi.xlsx")
    a = ap.parse_args(argv)
    if not a.mizan:
        a.mizan = [f"2025={BURASI / 'ornek_veri' / 'mizan_2025.csv'}", f"2026={BURASI / 'ornek_veri' / 'mizan_2026.csv'}"]
    mizanlar: "OrderedDict[str, Path]" = OrderedDict()
    for x in a.mizan:
        ad, _, yol = x.partition("=")
        if not yol:
            ad, yol = Path(x).stem, x
        mizanlar[ad] = Path(yol)
    gunler = {ad: para(g) for ad, g in zip(mizanlar, a.gun)}
    s = calistir(mizanlar, a.cikti, gunler)
    son = list(s["donemler"])[-1]
    r = s["donemler"][son]["rasyolar"]
    print(f"[OK] {len(mizanlar)} dönem · son dönem {son}")
    bicimler = {ad: b for _, ad, b, *_ in RASYOLAR}
    for ad in ("Cari oran", "Asit-test oranı", "Kaldıraç oranı", "Nakit dönüşüm süresi", "Net kâr marjı", "Özkaynak kârlılığı (ROE)"):
        v = r.get(ad)
        if v is not None:
            print(f"     {ad:<28} " + (f"%{v * 100:.1f}" if bicimler[ad] == "%" else f"{v:.2f}" + (" gün" if bicimler[ad] == "gün" else "")))
    for u in s["uyarilar"]:
        print(f"[!] {u}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
