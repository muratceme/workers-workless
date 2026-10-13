"""
İşe Alım Süreç Raporu — Workers / Workless kod bloğu
İnsan Kaynakları › İşe Alım Uzmanı

Aday takip listesinden açık pozisyonların aşama dağılımını, bekleme sürelerini, kaynak kanal kırılımını ve işe alım süresini raporlar.
  Aşama                = Başvuru · Ön Görüşme · Mülakat · Teklif (aktif aşamalar) · İşe Alındı · Olumsuz · Çekildi
  İşe alım süresi      = işe alınan adayın "İşe Alındı" aşama tarihi − başvuru tarihi (gün); pozisyon ve genel ortalama        (time-to-hire)
  Pozisyon açık gün    = (kapanış tarihi, yoksa bugün) − açılış tarihi;  kapanmış pozisyonda bu "doldurma süresi"dir           (time-to-fill)
  Aşamada bekleme      = bugün − adayın son aşama tarihi;  aktif aşamada --bekleme (7) günü aşan adaylar listelenir
  Kaynak dönüşümü %    = kaynaktan işe alınan ÷ kaynaktan gelen aday × 100;  mülakata ulaşma % = Mülakat, Teklif veya İşe Alındı aşamasındakiler ÷ aday
Rapor adayların yalnız şimdiki aşamasını bilir: olumsuz sonuçlanan adayın hangi aşamada elendiği dosyada yoksa huni analizi yapılmaz.
Aday verisi kişisel veridir; raporu yalnız işe alım sürecindeki yetkililerle paylaşın.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


ADAY = {"pozisyon": ("pozisyon", "ilan", "pozisyon adi"), "aday": ("aday", "aday adi", "ad soyad"), "kaynak": ("kaynak", "kaynak kanal", "kanal"), "basvuru": ("basvuru tarihi", "basvuru"), "asama": ("asama", "durum", "surec asamasi"),
        "tarih": ("asama tarihi", "son islem tarihi", "guncelleme tarihi", "durum tarihi")}
POZ = {"pozisyon": ADAY["pozisyon"], "departman": ("departman", "birim"), "acilis": ("acilis tarihi", "acilis", "talep tarihi"), "kapanis": ("kapanis tarihi", "kapanis"), "sorumlu": ("sorumlu", "ise alim uzmani")}
ASAMA = {"basvuru": "Başvuru", "yeni": "Başvuru", "on gorusme": "Ön Görüşme", "on eleme": "Ön Görüşme", "telefon gorusmesi": "Ön Görüşme", "mulakat": "Mülakat", "gorusme": "Mülakat", "teknik mulakat": "Mülakat", "ik mulakati": "Mülakat",
         "teklif": "Teklif", "teklif yapildi": "Teklif", "ise alindi": "İşe Alındı", "kabul": "İşe Alındı", "teklif kabul": "İşe Alındı", "basladi": "İşe Alındı", "olumsuz": "Olumsuz", "red": "Olumsuz", "reddedildi": "Olumsuz",
         "elendi": "Olumsuz", "cekildi": "Çekildi", "vazgecti": "Çekildi", "teklifi reddetti": "Çekildi"}
AKTIF = ("Başvuru", "Ön Görüşme", "Mülakat", "Teklif")
SIRA = AKTIF + ("İşe Alındı", "Olumsuz", "Çekildi")


def oku(aday_yol: Path, poz_yol: Path | None = None) -> tuple[list[dict], dict, list[dict]]:
    uy, adaylar, pozlar = [], [], {}
    for r in kayitlar(aday_yol, ADAY, ("pozisyon", "aday", "basvuru", "asama")):
        p, a = metin(r["pozisyon"]), metin(r["aday"])
        if not p and not a:
            continue
        b, asama, t = tarih(r["basvuru"]), ASAMA.get(katla(r["asama"])), tarih(r["tarih"])
        if not p or not a or b is None or asama is None or (t is None and metin(r["tarih"])) or (t and t < b):
            uy.append(uyari("Yüksek", "Aday satırı okunamadı", f"{p or '?'} · satır {r['_satir']}", f"Başvuru tarihi '{metin(r['basvuru'])}', aşama '{metin(r['asama'])}', aşama tarihi '{metin(r['tarih'])}'. "
                            f"Aşama şunlardan biri olmalı: {', '.join(SIRA)}; aşama tarihi başvurudan önce olamaz. Satır alınmadı."))
            continue
        adaylar.append({"pozisyon": p, "aday": a, "kaynak": metin(r["kaynak"]) or "Belirtilmemiş", "basvuru": b, "asama": asama, "tarih": t or b})
    if poz_yol:
        for r in kayitlar(poz_yol, POZ, ("pozisyon", "acilis")):
            p = metin(r["pozisyon"])
            if not p:
                continue
            a, k = tarih(r["acilis"]), tarih(r["kapanis"])
            if a is None or (k is None and metin(r["kapanis"])) or (k and k < a):
                uy.append(uyari("Orta", "Pozisyon satırı okunamadı", p, f"Açılış '{metin(r['acilis'])}', kapanış '{metin(r['kapanis'])}'. Pozisyonun açık gün sayısı hesaplanmadı."))
                continue
            pozlar[katla(p)] = {"pozisyon": p, "departman": metin(r["departman"]), "acilis": a, "kapanis": k, "sorumlu": metin(r["sorumlu"])}
    return adaylar, pozlar, uy


def ort(d: list[int]) -> Decimal | None:
    return yuvarla(Decimal(sum(d)) / len(d), "0.1") if d else None


def rapor(adaylar: list[dict], pozlar: dict | None, bugun: date, bekleme: int = 7) -> dict:
    if not adaylar:
        raise ValueError("Raporlanacak aday yok.")
    pozlar, uy, grup, kaynak, bekleyen = pozlar or {}, [], {}, {}, []
    for a in adaylar:
        if a["tarih"] > bugun:
            uy.append(uyari("Orta", "İleri tarihli aşama", f"{a['pozisyon']} · {a['aday']}", f"Aşama tarihi {a['tarih']:%d.%m.%Y} rapor tarihinden sonra; bekleme süresi 0 sayıldı."))
        pk = katla(a["pozisyon"])
        g = grup.setdefault(pk, {"pozisyon": a["pozisyon"], "aday": 0, "sureler": [], **{s: 0 for s in SIRA}})
        g["aday"] += 1
        g[a["asama"]] += 1
        k = kaynak.setdefault(katla(a["kaynak"]), {"kaynak": a["kaynak"], "aday": 0, "mulakat": 0, "alinan": 0, "sureler": []})
        k["aday"] += 1
        k["mulakat"] += a["asama"] in ("Mülakat", "Teklif", "İşe Alındı")
        if a["asama"] == "İşe Alındı":
            sure = (a["tarih"] - a["basvuru"]).days
            g["sureler"].append(sure)
            k["sureler"].append(sure)
            k["alinan"] += 1
        elif a["asama"] in AKTIF:
            gun = max((bugun - a["tarih"]).days, 0)
            if gun > bekleme:
                bekleyen.append({"pozisyon": a["pozisyon"], "aday": a["aday"], "asama": a["asama"], "tarih": a["tarih"], "gun": gun, "kaynak": a["kaynak"]})
    liste, tum_sure = [], []
    for pk, g in grup.items():
        p, sureler = pozlar.get(pk), g.pop("sureler")
        tum_sure += sureler
        aktif, bek = sum(g[s] for s in AKTIF), sum(1 for b in bekleyen if katla(b["pozisyon"]) == pk)
        durum = "Kapandı" if p and p["kapanis"] else "Açık"
        x = g | {"departman": p["departman"] if p else "", "sorumlu": p["sorumlu"] if p else "", "acilis": p["acilis"] if p else None, "kapanis": p["kapanis"] if p else None, "durum": durum,
                 "acik_gun": ((p["kapanis"] or bugun) - p["acilis"]).days if p else None, "aktif": aktif, "bekleyen": bek, "ise_alim_suresi": ort(sureler)}
        liste.append(x)
        if durum == "Açık" and not aktif:
            uy.append(uyari("Yüksek", "Aktif adayı olmayan açık pozisyon", g["pozisyon"], f"{g['aday']} adayın hiçbiri aktif aşamada değil" + (f"; pozisyon {x['acik_gun']} gündür açık." if p else ".") + " Aday havuzunu yenileyin veya pozisyonu kapatın."))
    for pk, p in pozlar.items():
        if pk not in grup:
            liste.append({"pozisyon": p["pozisyon"], "aday": 0, **{s: 0 for s in SIRA}, "departman": p["departman"], "sorumlu": p["sorumlu"], "acilis": p["acilis"], "kapanis": p["kapanis"], "durum": "Kapandı" if p["kapanis"] else "Açık",
                          "acik_gun": ((p["kapanis"] or bugun) - p["acilis"]).days, "aktif": 0, "bekleyen": 0, "ise_alim_suresi": None})
            if not p["kapanis"]:
                uy.append(uyari("Yüksek", "Adayı olmayan açık pozisyon", p["pozisyon"], f"Pozisyon {liste[-1]['acik_gun']} gündür açık ama aday listesinde kaydı yok."))
    liste.sort(key=lambda x: (x["durum"] != "Açık", -(x["acik_gun"] or 0), x["pozisyon"]))
    bekleyen.sort(key=lambda b: (-b["gun"], b["pozisyon"], b["aday"]))
    for b in bekleyen:
        uy.append(uyari("Orta", "Aşamada bekleyen aday", f"{b['pozisyon']} · {b['aday']}", f"{b['asama']} aşamasında {b['gun']} gündür işlem yok (son işlem {b['tarih']:%d.%m.%Y})."))
    kaynaklar = sorted(({"kaynak": k["kaynak"], "aday": k["aday"], "mulakat": k["mulakat"], "alinan": k["alinan"], "mulakat_y": yuvarla(Decimal(k["mulakat"]) / k["aday"] * 100), "donusum": yuvarla(Decimal(k["alinan"]) / k["aday"] * 100),
                         "sure": ort(k["sureler"])} for k in kaynak.values()), key=lambda k: (-k["aday"], k["kaynak"]))
    kapali = [x["acik_gun"] for x in liste if x["durum"] == "Kapandı" and x["acik_gun"] is not None]
    genel = {"pozisyon": len(liste), "acik": sum(1 for x in liste if x["durum"] == "Açık"), "aday": len(adaylar), "aktif": sum(x["aktif"] for x in liste), "alinan": sum(x["İşe Alındı"] for x in liste), "ise_alim_suresi": ort(tum_sure),
             "doldurma_suresi": ort(kapali), "bekleyen": len(bekleyen), **{s: sum(x[s] for x in liste) for s in SIRA}}
    return {"pozisyonlar": liste, "kaynaklar": kaynaklar, "bekleyenler": bekleyen, "genel": genel, "bugun": bugun, "bekleme": bekleme, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb, g = Workbook(), r["genel"]
    ws = wb.active
    ws.title = "Pozisyon Özeti"
    _baslik(ws, ["Pozisyon", "Departman", "Durum", "Açılış", "Kapanış", "Açık Gün", "Aday"] + list(SIRA) + ["Aktif Aday", "Bekleyen", "Ort. İşe Alım Süresi (gün)", "Sorumlu"], (30, 18, 10, 12, 12, 9, 7) + (10,) * len(SIRA) + (9, 9, 12, 20), "B2")
    for x in r["pozisyonlar"]:
        _satir(ws, [x["pozisyon"], x["departman"], x["durum"], x["acilis"], x["kapanis"], x["acik_gun"], x["aday"]] + [x[s] or None for s in SIRA] + [x["aktif"] or None, x["bekleyen"] or None, x["ise_alim_suresi"], x["sorumlu"]], tarih_s=(4, 5),
               dolgu={3: YESIL if x["durum"] == "Kapandı" else None, 9 + len(SIRA): SARI if x["bekleyen"] else None, 8 + len(SIRA): KIRMIZI if x["durum"] == "Açık" and not x["aktif"] else None})
    n = _satir(ws, ["Genel", "", f"{g['acik']} açık", None, None, g["doldurma_suresi"], g["aday"]] + [g[s] or None for s in SIRA] + [g["aktif"] or None, g["bekleyen"] or None, g["ise_alim_suresi"], ""])
    for c in ws[n]:
        c.font = Font(bold=True)
    ws.append([])
    ws.append([f"Rapor tarihi {r['bugun']:%d.%m.%Y}. İşe alım süresi = işe alındı tarihi − başvuru tarihi. Genel satırındaki 'Açık Gün' kapanmış pozisyonların ortalama doldurma süresidir. Bekleyen: aktif aşamada {r['bekleme']} günden uzun işlem görmeyen aday."])
    ws = wb.create_sheet("Kaynak Analizi")
    _baslik(ws, ["Kaynak", "Aday", "Mülakata Ulaşan", "Mülakata Ulaşma %", "İşe Alınan", "Dönüşüm %", "Ort. İşe Alım Süresi (gün)"], (26, 8, 12, 12, 10, 11, 14))
    for k in r["kaynaklar"]:
        _satir(ws, [k["kaynak"], k["aday"], k["mulakat"] or None, k["mulakat_y"], k["alinan"] or None, k["donusum"], k["sure"]], para_s=(4, 6))
    ws = wb.create_sheet("Bekleyen Adaylar")
    _baslik(ws, ["Pozisyon", "Aday", "Aşama", "Son İşlem", "Bekleme (gün)", "Kaynak", "Aksiyon (siz doldurun)"], (30, 26, 14, 12, 11, 20, 30))
    for b in r["bekleyenler"]:
        _satir(ws, [b["pozisyon"], b["aday"], b["asama"], b["tarih"], b["gun"], b["kaynak"], ""], tarih_s=(4,), dolgu={5: KIRMIZI if b["gun"] > 2 * r["bekleme"] else SARI})
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(aday_yol: Path, cikti: Path, bugun: date, poz_yol: Path | None = None, bekleme: int = 7) -> dict:
    adaylar, pozlar, uy = oku(aday_yol, poz_yol)
    r = rapor(adaylar, pozlar, bugun, bekleme)
    uy += r["uyarilar"]
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Aday takip listesinden işe alım süreç raporu üretir.")
    p.add_argument("--adaylar", type=Path, default=ORNEK / "adaylar.csv", help="Pozisyon, Aday, Kaynak, Başvuru Tarihi, Aşama, Aşama Tarihi")
    p.add_argument("--pozisyonlar", type=Path, help="İsteğe bağlı: Pozisyon, Departman, Açılış Tarihi, Kapanış Tarihi, Sorumlu")
    p.add_argument("--tarih", help="Rapor tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--bekleme", type=int, default=7, help="Aktif aşamada bu günden uzun bekleyen adaylar listelenir (varsayılan 7)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "ise_alim_surec_raporu.xlsx")
    a = p.parse_args(argv)
    ornek = a.adaylar == ORNEK / "adaylar.csv"
    poz = a.pozisyonlar or (ORNEK / "pozisyonlar.csv" if ornek else None)
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 13) if ornek else date.today()
    if bugun is None or a.bekleme < 0:
        print("[X] --tarih GG.AA.YYYY biçiminde, --bekleme sıfır veya pozitif olmalı")
        return 2
    for y in (a.adaylar, poz):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.adaylar, a.cikti, bugun, poz, a.bekleme)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    g = r["genel"]
    print(f"[OK] {g['pozisyon']} pozisyon ({g['acik']} açık) · {g['aday']} aday · {g['aktif']} aktif · {g['alinan']} işe alınan · ortalama işe alım süresi {g['ise_alim_suresi'] if g['ise_alim_suresi'] is not None else '-'} gün · {g['bekleyen']} bekleyen")
    for b in r["bekleyenler"][:10]:
        print(f"[!] {b['pozisyon']} · {b['aday']}: {b['asama']} aşamasında {b['gun']} gün")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
