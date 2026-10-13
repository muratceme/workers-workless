import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
B = date(2026, 10, 13)
g = lambda ay, gun: date(2026, ay, gun)   # noqa: E731


def a(poz, aday, kaynak, basvuru, asama, tarih=None):
    return {"pozisyon": poz, "aday": aday, "kaynak": kaynak, "basvuru": basvuru, "asama": asama, "tarih": tarih or basvuru}


ADAY = lambda: [a("Satış", "A1", "LinkedIn", g(9, 1), "İşe Alındı", g(10, 1)), a("Satış", "A2", "Kariyer", g(9, 5), "Olumsuz", g(9, 20)), a("Satış", "A3", "linkedin", g(9, 10), "Mülakat", g(10, 1)),   # noqa: E731
                a("SATIŞ", "A4", "Referans", g(10, 1), "Başvuru"), a("Muhasebe", "B1", "Kariyer", g(9, 20), "Teklif", g(10, 10)), a("Muhasebe", "B2", "Kariyer", g(9, 25), "İşe Alındı", g(10, 5))]
POZ = lambda: {"satis": {"pozisyon": "Satış", "departman": "S", "acilis": g(8, 25), "kapanis": None, "sorumlu": "X"}, "muhasebe": {"pozisyon": "Muhasebe", "departman": "M", "acilis": g(9, 15), "kapanis": g(10, 5), "sorumlu": "Y"}}   # noqa: E731


class RaporTesti(unittest.TestCase):
    def test_pozisyon_ozeti(self):
        r = main.rapor(ADAY(), POZ(), B)
        self.assertEqual([(x["pozisyon"], x["durum"], x["acik_gun"], x["aday"], x["Başvuru"], x["Mülakat"], x["Teklif"], x["İşe Alındı"], x["Olumsuz"], x["aktif"], x["bekleyen"], x["ise_alim_suresi"]) for x in r["pozisyonlar"]],
                         [("Satış", "Açık", 49, 4, 1, 1, 0, 1, 1, 2, 2, D("30.0")), ("Muhasebe", "Kapandı", 20, 2, 0, 0, 1, 1, 0, 1, 0, D("10.0"))])
        gn = r["genel"]
        self.assertEqual((gn["pozisyon"], gn["acik"], gn["aday"], gn["aktif"], gn["alinan"], gn["ise_alim_suresi"], gn["doldurma_suresi"], gn["bekleyen"], gn["Olumsuz"]), (2, 1, 6, 3, 2, D("20.0"), D("20.0"), 2, 1))

    def test_kaynak_ve_bekleyenler(self):
        r = main.rapor(ADAY(), POZ(), B)
        self.assertEqual([(k["kaynak"], k["aday"], k["mulakat"], k["mulakat_y"], k["alinan"], k["donusum"], k["sure"]) for k in r["kaynaklar"]],
                         [("Kariyer", 3, 2, D("66.67"), 1, D("33.33"), D("10.0")), ("LinkedIn", 2, 2, D("100.00"), 1, D("50.00"), D("30.0")), ("Referans", 1, 0, D("0.00"), 0, D("0.00"), None)])
        self.assertEqual([(b["aday"], b["asama"], b["gun"]) for b in r["bekleyenler"]], [("A4", "Başvuru", 12), ("A3", "Mülakat", 12)])
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]), [("Aşamada bekleyen aday", "SATIŞ · A4"), ("Aşamada bekleyen aday", "Satış · A3")])
        self.assertEqual(([b["aday"] for b in main.rapor(ADAY(), POZ(), B, bekleme=2)["bekleyenler"]], main.rapor(ADAY(), POZ(), B, bekleme=12)["bekleyenler"]), (["A4", "A3", "B1"], []))

    def test_pozisyonsuz_ve_uyarilar(self):
        r = main.rapor(ADAY(), None, B)
        self.assertEqual([(x["pozisyon"], x["durum"], x["acik_gun"], x["acilis"]) for x in r["pozisyonlar"]], [("Muhasebe", "Açık", None, None), ("Satış", "Açık", None, None)])
        poz = POZ() | {"depo": {"pozisyon": "Depo", "departman": "L", "acilis": g(9, 1), "kapanis": None, "sorumlu": ""}, "eski": {"pozisyon": "Eski", "departman": "", "acilis": g(1, 1), "kapanis": g(2, 1), "sorumlu": ""}}
        r = main.rapor([a("Satış", "A2", "K", g(9, 5), "Olumsuz", g(9, 20)), a("Satış", "A9", "K", g(10, 10), "Mülakat", g(10, 20))], poz, B)
        self.assertEqual([(x["pozisyon"], x["durum"], x["acik_gun"], x["aday"]) for x in r["pozisyonlar"]], [("Satış", "Açık", 49, 2), ("Depo", "Açık", 42, 0), ("Eski", "Kapandı", 31, 0), ("Muhasebe", "Kapandı", 20, 0)])
        self.assertEqual((sorted((u["tur"], u["kim"]) for u in r["uyarilar"]), r["bekleyenler"], r["genel"]["ise_alim_suresi"], r["genel"]["doldurma_suresi"]),
                         ([("Adayı olmayan açık pozisyon", "Depo"), ("İleri tarihli aşama", "Satış · A9")], [], None, D("25.5")))
        r = main.rapor([a("Satış", "A2", "K", g(9, 5), "Olumsuz", g(9, 20))], POZ(), B)
        self.assertEqual([u["tur"] for u in r["uyarilar"]], ["Aktif adayı olmayan açık pozisyon"])
        with self.assertRaises(ValueError):
            main.rapor([], None, B)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "i.xlsx"
            r = main.calistir(ORNEK / "adaylar.csv", cikti, B, ORNEK / "pozisyonlar.csv")
            self.assertEqual([(x["pozisyon"], x["durum"], x["acik_gun"], x["aday"], x["aktif"], x["ise_alim_suresi"]) for x in r["pozisyonlar"]],
                             [("Satış Uzmanı", "Açık", 49, 4, 2, D("30.0")), ("Depo Sorumlusu", "Açık", 42, 0, 0, None), ("Muhasebe Uzmanı", "Kapandı", 20, 2, 1, D("10.0"))])
            self.assertEqual(([(u["onem"], u["tur"]) for u in r["uyarilar"]], r["genel"]["ise_alim_suresi"]), ([("Yüksek", "Adayı olmayan açık pozisyon"), ("Orta", "Aşamada bekleyen aday"), ("Orta", "Aşamada bekleyen aday")], D("20.0")))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Pozisyon Özeti", "Kaynak Analizi", "Bekleyen Adaylar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            ay, py = Path(tmp) / "a.csv", Path(tmp) / "p.csv"
            ay.write_text("İlan;Ad Soyad;Kanal;Başvuru;Durum;Durum Tarihi\nP;A;;01.09.2026;ön eleme;05.09.2026\nP;B;K;01.09.2026;Bekliyor;\nP;C;K;x;Mülakat;\nP;D;K;10.09.2026;TEKLİF;01.09.2026\nP;E;K;10.09.2026;Reddedildi;\n", encoding="utf-8")
            py.write_text("Pozisyon;Açılış;Kapanış\nP;01.08.2026;\nQ;x;\nR;01.08.2026;01.07.2026\n", encoding="utf-8")
            adaylar, pozlar, uy = main.oku(ay, py)
            self.assertEqual(([(x["aday"], x["kaynak"], x["asama"], x["tarih"]) for x in adaylar], list(pozlar), sorted(u["tur"] for u in uy)),
                             ([("A", "Belirtilmemiş", "Ön Görüşme", g(9, 5)), ("E", "K", "Olumsuz", g(9, 10))], ["p"], ["Aday satırı okunamadı"] * 3 + ["Pozisyon satırı okunamadı"] * 2))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--adaylar", str(ay), "--cikti", c]
            self.assertEqual(main.main(arg + ["--pozisyonlar", str(py), "--tarih", "13.10.2026", "--bekleme", "14"]), 0)
            self.assertEqual(main.main(arg), 0)
            for ek in (["--tarih", "x"], ["--bekleme", "-1"]):
                self.assertEqual(main.main(arg + ek), 2)
            self.assertEqual(main.main(arg + ["--pozisyonlar", str(Path(tmp) / "yok.csv")]), 1)
            ay.write_text("Pozisyon;Aday\nP;A\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
