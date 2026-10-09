"""
BT Maliyet Raporu — Workers / Workless kod bloğu
Bilgi Teknolojileri › BT Müdürü

Lisans, bulut, donanım, hizmet ve iletişim harcamalarını birim ve kalem bazında raporlar:
  - Kategori: dosyada yoksa açıklamadan (bulut → iletişim → donanım → lisans → hizmet sırasıyla anahtar kelime).
  - TL karşılığı: tutar × kur; kur yoksa dosyadaki aynı dövizin en yakın tarihli kuru kullanılır (uyarıyla).
  - Tahakkuk görünümü: "Yıllık" ödemeler başlangıç ayından itibaren 12 aya yayılır; rapor yılını aşan kısım peşin
    ödenmiş gider olarak ayrı gösterilir. "Aylık" ve "Tek seferlik" ödendiği aya yazılır.
  - Bütçe karşılaştırması: kategori yıllık bütçesi ↔ dönem tahakkuku ve yıl sonu tahmini (aylık ortalama × 12;
    yıllık ödemeler kendi tutarıyla).
  - Kişi başı BT maliyeti (birim çalışan sayıları verilirse; "Genel" birim çalışan sayısına göre dağıtılmaz, ayrı gösterilir).
  - Lisans kullanımı: atıl lisans ve maliyeti, kullanım oranı < %80.
  - Kontroller: aylık harcaması önceki 3 ay ortalamasından %30 fazla artan kalem (ör. bulut), farklı birimlerin aynı
    ay aynı tedarikçiden aynı aboneliği alması (olası mükerrer), yaklaşan yenileme / sözleşme bitişi (--yenileme-gun),
    kuru olmayan döviz harcaması.
Rapor: özet, kategori × ay, birim × kategori, tedarikçiler, bütçe, lisans kullanımı, yenileme takvimi, uyarılar.
İnternete bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: 49 harcama, Ocak–Eylül 2026
    python main.py --harcamalar h.xlsx --lisanslar l.xlsx --butce b.csv --birimler birimler.csv --yil 2026 --bugun 09.10.2026
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
K2 = Decimal("0.01")
SIFIR = Decimal(0)
AYLAR = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]
KATEGORILER = ["Lisans", "Bulut", "Donanım", "Hizmet", "İletişim", "Diğer"]
KURALLAR = [("Bulut", ("bulut", "cloud", "aws", "azure", "google cloud", "sunucu kiralama", "hosting", "depolama")),
            ("İletişim", ("internet", "hat", "gsm", "telekom", "mpls", "data hatti", "telefon")),
            ("Donanım", ("bilgisayar", "laptop", "dizustu", "sunucu", "disk", "bellek", "yazici", "donanim", "monitor", "switch", "telefon cihazi")),
            ("Lisans", ("lisans", "abonelik", "license", "saas", "365", "subscription")),
            ("Hizmet", ("hizmet", "destek", "danismanlik", "test", "bakim", "egitim", "kurulum", "dis kaynak"))]

SUTUNLAR = {"tarih": ("tarih", "fatura tarihi"), "tedarikci": ("tedarikci", "satici", "firma"), "kalem": ("kalem", "aciklama", "urun hizmet"),
            "kategori": ("kategori",), "birim": ("birim", "departman", "maliyet merkezi"), "tutar": ("tutar",), "doviz": ("doviz", "para birimi"),
            "kur": ("kur",), "tip": ("donem tipi", "odeme tipi", "periyot"), "bitis": ("sozlesme bitis", "bitis tarihi", "yenileme tarihi")}
LISANS_SUTUNLARI = {"urun": ("urun", "lisans"), "toplam": ("toplam lisans", "satin alinan"), "kullanilan": ("kullanilan lisans", "atanmis", "kullanilan"),
                    "fiyat": ("yillik birim fiyat", "birim fiyat"), "doviz": ("doviz", "para birimi"), "yenileme": ("yenileme tarihi", "bitis")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tl(x: Decimal | None) -> str:
    return "—" if x is None else f"{x.quantize(K2, ROUND_HALF_UP):,.0f}".replace(",", ".")


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


def kategori_bul(kalem: str, verilen: str = "") -> str:
    v = katla(verilen)
    for k in KATEGORILER:
        if v and v == katla(k):
            return k
    t = f" {katla(kalem)} "
    for kat, kelimeler in KURALLAR:
        if any(f" {w} " in t or (len(w) > 4 and f" {w}" in t) for w in kelimeler):
            return kat
    return "Diğer"


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Harcama:
    tarih: date
    tedarikci: str
    kalem: str
    kategori: str
    birim: str
    tutar: Decimal
    doviz: str
    kur: Decimal | None
    tip: str             # aylik | yillik | tek
    bitis: date | None
    satir: int
    notlar: list = field(default_factory=list)

    @property
    def tl(self) -> Decimal:
        return (self.tutar * (self.kur or Decimal(1))).quantize(K2, ROUND_HALF_UP) if self.doviz != "TRY" else self.tutar


def oku(yol: Path) -> tuple[list[Harcama], list[dict]]:
    sonuc, uy = [], []
    for r in kayitlar(yol, SUTUNLAR, ("tarih", "kalem", "tutar")):
        t, m = tarih(r.get("tarih")), para(r.get("tutar"))
        if not t or m is None:
            if any(r.get(k) for k in ("kalem", "tutar")):
                uy.append({"onem": "Orta", "tur": "Okunamayan satır", "kim": metin(r.get("kalem")), "aciklama": f"Satır {r['_satir']}: tarih veya tutar okunamadı"})
            continue
        tip = katla(r.get("tip"))
        tip = "yillik" if tip.startswith(("yil", "annual")) else "tek" if tip.startswith(("tek", "bir kez", "one")) else "aylik"
        dv = (metin(r.get("doviz")) or "TRY").upper().replace("TL", "TRY")
        sonuc.append(Harcama(t, metin(r.get("tedarikci")), metin(r.get("kalem")), kategori_bul(metin(r.get("kalem")), metin(r.get("kategori"))),
                             metin(r.get("birim")) or "Genel", m, dv, para(r.get("kur")), tip, tarih(r.get("bitis")), r["_satir"]))
    # Eksik kurlar: aynı dövizin en yakın tarihli kuru
    for h in sonuc:
        if h.doviz != "TRY" and not h.kur:
            adaylar = [x for x in sonuc if x.doviz == h.doviz and x.kur]
            if adaylar:
                y = min(adaylar, key=lambda x: abs((x.tarih - h.tarih).days))
                h.kur = y.kur
                h.notlar.append(f"Kur yoktu; {y.tarih:%d.%m.%Y} kuru ({y.kur}) kullanıldı")
                uy.append({"onem": "Orta", "tur": "Kur tahmini", "kim": h.kalem, "aciklama": f"{h.doviz} kuru verilmemiş; dosyadaki en yakın tarihli kur "
                           f"({y.tarih:%d.%m.%Y}: {y.kur}) kullanıldı. Fatura kurunu girin"})
            else:
                uy.append({"onem": "Yüksek", "tur": "Kur yok", "kim": h.kalem, "aciklama": f"{h.doviz} harcamasının kuru yok; TL toplamlara 1 kurla girdi"})
    return sonuc, uy


# ----------------------------------------------------------------------------
# Analiz
# ----------------------------------------------------------------------------

def tahakkuk(h: Harcama, yil: int) -> tuple[dict[int, Decimal], Decimal]:
    """ay → tutar (rapor yılı içinde), gelecek yıla kalan peşin kısım"""
    if h.tip != "yillik":
        return ({h.tarih.month: h.tl} if h.tarih.year == yil else {}), SIFIR
    aylik = (h.tl / 12)
    d, kalan = {}, SIFIR
    for i in range(12):
        ay = h.tarih.month + i
        y = h.tarih.year + (ay - 1) // 12
        a = (ay - 1) % 12 + 1
        if y == yil:
            d[a] = d.get(a, SIFIR) + aylik
        elif y > yil:
            kalan += aylik
    return d, kalan


def analiz_et(harcamalar: list[Harcama], yil: int, bugun: date, lisanslar: list[dict], butce: dict[str, Decimal], birimler: dict[str, int],
              yenileme_gun: int, artis_esik: Decimal) -> dict:
    uy = []
    son_ay = max((h.tarih.month for h in harcamalar if h.tarih.year == yil), default=12)
    kat_ay = defaultdict(lambda: defaultdict(lambda: SIFIR))
    birim_kat = defaultdict(lambda: defaultdict(lambda: SIFIR))
    tedarikci = defaultdict(lambda: SIFIR)
    pesin = defaultdict(lambda: SIFIR)
    nakit = SIFIR
    for h in harcamalar:
        d, kalan = tahakkuk(h, yil)
        for a, v in d.items():
            if a <= son_ay:
                kat_ay[h.kategori][a] += v
                birim_kat[h.birim][h.kategori] += v
            else:
                kalan += v
        if h.tarih.year == yil:
            tedarikci[h.tedarikci] += h.tl
            nakit += h.tl
        if kalan:
            pesin[h.kategori] += kalan
    # Aylık artış (aylık kalemler, kalem + tedarikçi bazında)
    seri = defaultdict(dict)
    for h in harcamalar:
        if h.tip == "aylik" and h.tarih.year == yil:
            anahtar = (h.tedarikci, katla(h.kalem), h.birim)
            seri[anahtar][h.tarih.month] = seri[anahtar].get(h.tarih.month, SIFIR) + h.tl
    for (ted, _, birim), s in seri.items():
        for a in sorted(s):
            onceki = [s[x] for x in (a - 3, a - 2, a - 1) if x in s]
            if len(onceki) == 3:
                ort = sum(onceki) / 3
                if ort and (s[a] - ort) / ort > artis_esik:
                    uy.append({"onem": "Orta", "tur": "Harcama artışı", "kim": f"{ted} · {birim}", "aciklama": f"{AYLAR[a - 1]}: {tl(s[a])} TL, önceki 3 ay "
                               f"ortalaması {tl(ort)} TL (%{(s[a] - ort) / ort * 100:.0f}); kullanım veya fiyat değişikliğini kontrol edin"})
    # Olası mükerrer abonelik
    grup = defaultdict(set)
    for h in harcamalar:
        if h.tip in ("aylik", "yillik") and h.kategori in ("Lisans", "Bulut"):
            grup[(katla(h.tedarikci), katla(h.kalem), h.tarih.year, h.tarih.month)].add(h.birim)
    goruldu = set()
    for (ted, kalem, y, a), bs in sorted(grup.items()):
        if len(bs) > 1 and (ted, kalem) not in goruldu:
            goruldu.add((ted, kalem))
            ornek = next(h for h in harcamalar if katla(h.tedarikci) == ted and katla(h.kalem) == kalem)
            uy.append({"onem": "Orta", "tur": "Olası mükerrer abonelik", "kim": ornek.tedarikci, "aciklama": f"'{ornek.kalem}' {AYLAR[a - 1]} {y} itibarıyla "
                       f"{len(bs)} birim tarafından ayrı ayrı alınıyor ({', '.join(sorted(bs))}); kurumsal / ekip planında birleştirilebilir"})
    # Yenilemeler
    yenileme = []
    for h in harcamalar:
        if h.bitis and 0 <= (h.bitis - bugun).days <= yenileme_gun:
            yenileme.append({"ne": h.kalem, "tedarikci": h.tedarikci, "tarih": h.bitis, "tutar": h.tl, "kaynak": "Sözleşme"})
    for l in lisanslar:
        if l["yenileme"] and 0 <= (l["yenileme"] - bugun).days <= yenileme_gun:
            yenileme.append({"ne": l["urun"], "tedarikci": "", "tarih": l["yenileme"], "tutar": None, "kaynak": "Lisans"})
    tekil = {}
    for y in sorted(yenileme, key=lambda y: (y["tarih"], y["kaynak"] != "Sözleşme")):
        tekil.setdefault(((katla(y["ne"]).split() or [""])[0], y["tarih"]), y)
    yenileme = list(tekil.values())
    for y in yenileme:
        uy.append({"onem": "Bilgi", "tur": "Yaklaşan yenileme", "kim": y["ne"], "aciklama": f"{y['tarih']:%d.%m.%Y} ({(y['tarih'] - bugun).days} gün); kullanım "
                   "ve ihtiyaç gözden geçirilip pazarlık yapılmalı"})
    # Lisans kullanımı
    lk = []
    for l in lisanslar:
        if not l["toplam"]:
            continue
        atil = max(l["toplam"] - l["kullanilan"], 0)
        oran = l["kullanilan"] / l["toplam"]
        lk.append({**l, "atil": atil, "oran": oran, "atil_maliyet": atil * l["fiyat"] if l["fiyat"] is not None else None})
        if oran < Decimal("0.8"):
            uy.append({"onem": "Orta", "tur": "Atıl lisans", "kim": l["urun"], "aciklama": f"{l['kullanilan']:g} / {l['toplam']:g} kullanımda (%{oran * 100:.0f}); "
                       f"{atil:g} atıl lisansın yıllık maliyeti {tl(atil * l['fiyat']) if l['fiyat'] is not None else '—'} {l['doviz']}. Yenilemede azaltılabilir"})
        if l["kullanilan"] > l["toplam"]:
            uy.append({"onem": "Yüksek", "tur": "Lisans aşımı", "kim": l["urun"], "aciklama": f"{l['kullanilan']:g} kullanıcı, {l['toplam']:g} lisans; uyum riski"})
    # Bütçe
    bt = []
    for kat in KATEGORILER:
        donem = sum(kat_ay.get(kat, {}).values(), SIFIR)
        if not donem and kat not in butce:
            continue
        aylik_ort = sum((v for h in harcamalar if h.kategori == kat and h.tip == "aylik" and h.tarih.year == yil for v in [h.tl]), SIFIR) / son_ay
        yillik = sum((h.tl for h in harcamalar if h.kategori == kat and h.tip == "yillik" and h.tarih.year == yil), SIFIR)
        tek = sum((h.tl for h in harcamalar if h.kategori == kat and h.tip == "tek" and h.tarih.year == yil), SIFIR)
        tahmin = aylik_ort * 12 + yillik + tek
        b = butce.get(kat)
        bt.append({"kategori": kat, "donem": donem, "butce": b, "oransal": b * son_ay / 12 if b is not None else None, "tahmin": tahmin,
                   "fark": tahmin - b if b is not None else None})
        if b is not None and tahmin > b:
            uy.append({"onem": "Orta" if tahmin <= b * Decimal("1.1") else "Yüksek", "tur": "Bütçe aşım tahmini", "kim": kat,
                       "aciklama": f"Yıl sonu tahmini {tl(tahmin)} TL, bütçe {tl(b)} TL (aşım {tl(tahmin - b)} TL)"})
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    uy.sort(key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))
    return {"harcamalar": harcamalar, "kat_ay": kat_ay, "birim_kat": birim_kat, "tedarikci": tedarikci, "pesin": pesin, "nakit": nakit, "son_ay": son_ay,
            "yenileme": yenileme, "lisans": lk, "butce": bt, "birimler": birimler, "uyarilar": uy, "yil": yil, "bugun": bugun}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
UST = Alignment(vertical="top", wrap_text=True)
PF = "#,##0"


def _baslik(ws, basliklar, genislik=()):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "B2"


def _f(x):
    return None if x is None else float(x)


def rapor_yaz(cikti: Path, s: dict) -> None:
    n = s["son_ay"]
    wb = Workbook()
    oz = wb.active
    oz.title = "Özet"
    _baslik(oz, ["Gösterge", "Tutar (TL)", "Açıklama"], (34, 16, 70))
    toplam = sum((sum(v.values(), SIFIR) for v in s["kat_ay"].values()), SIFIR)
    for a, v, ac in [("Tahakkuk toplamı", toplam, f"{s['yil']} Ocak – {AYLAR[n - 1]}; yıllık ödemeler aylara yayılmış"),
                     ("Ödenen (nakit) toplam", s["nakit"], "Fatura tarihine göre"),
                     ("Peşin ödenmiş, sonraki aylara ait", sum(s["pesin"].values(), SIFIR), "Yıllık aboneliklerin rapor dönemi sonrasına düşen kısmı"),
                     ("Aylık ortalama (tahakkuk)", toplam / n, ""),
                     ("Döviz cinsinden harcama payı", None, "")]:
        oz.append([a, _f(v), ac])
        oz.cell(oz.max_row, 2).number_format = PF
    dv = sum((h.tl for h in s["harcamalar"] if h.doviz != "TRY" and h.tarih.year == s["yil"]), SIFIR)
    oz.cell(oz.max_row, 2).value = float(dv / s["nakit"]) if s["nakit"] else None
    oz.cell(oz.max_row, 2).number_format = "0.0%"
    oz.append([])
    _baslik(oz, ["Birim", "Çalışan", "BT Maliyeti (tahakkuk)", "Kişi Başı"], ())
    for b, kats in sorted(s["birim_kat"].items()):
        t = sum(kats.values(), SIFIR)
        c = s["birimler"].get(b)
        oz.append([b, c, float(t), float(t / c) if c else None])
        for j in (3, 4):
            oz.cell(oz.max_row, j).number_format = PF

    ka = wb.create_sheet("Kategori × Ay")
    _baslik(ka, ["Kategori"] + AYLAR[:n] + ["Toplam", "Pay"], [16] + [12] * n + [14, 8])
    for kat in KATEGORILER:
        if kat in s["kat_ay"]:
            v = [s["kat_ay"][kat].get(a, SIFIR) for a in range(1, n + 1)]
            ka.append([kat] + [float(x) for x in v] + [float(sum(v)), float(sum(v) / toplam) if toplam else None])
            for j in range(2, n + 3):
                ka.cell(ka.max_row, j).number_format = PF
            ka.cell(ka.max_row, n + 3).number_format = "0.0%"
    satir_sayisi = ka.max_row
    ka.append(["Toplam"] + [float(sum((s["kat_ay"][k].get(a, SIFIR) for k in s["kat_ay"]), SIFIR)) for a in range(1, n + 1)] + [float(toplam)])
    for j in range(2, n + 3):
        ka.cell(ka.max_row, j).number_format = PF
        ka.cell(ka.max_row, j).font = Font(bold=True)
    g = BarChart()
    g.type, g.grouping, g.overlap = "col", "stacked", 100
    g.title, g.height, g.width = "Aylık BT maliyeti (tahakkuk)", 8, 20
    g.add_data(Reference(ka, min_col=1, max_col=n + 1, min_row=2, max_row=satir_sayisi), titles_from_data=True, from_rows=True)
    g.set_categories(Reference(ka, min_col=2, max_col=n + 1, min_row=1))
    ka.add_chart(g, f"A{ka.max_row + 3}")

    bk = wb.create_sheet("Birim × Kategori")
    kats = [k for k in KATEGORILER if k in s["kat_ay"]]
    _baslik(bk, ["Birim"] + kats + ["Toplam"], [16] + [13] * (len(kats) + 1))
    for b, kv in sorted(s["birim_kat"].items()):
        bk.append([b] + [float(kv.get(k, SIFIR)) for k in kats] + [float(sum(kv.values(), SIFIR))])
        for j in range(2, len(kats) + 3):
            bk.cell(bk.max_row, j).number_format = PF

    td = wb.create_sheet("Tedarikçiler")
    _baslik(td, ["Tedarikçi", "Ödenen (TL)", "Pay", "Kümülatif Pay"], (30, 15, 8, 12))
    kum = SIFIR
    for t, v in sorted(s["tedarikci"].items(), key=lambda i: -i[1]):
        kum += v
        td.append([t, float(v), float(v / s["nakit"]), float(kum / s["nakit"])])
        td.cell(td.max_row, 2).number_format = PF
        td.cell(td.max_row, 3).number_format = td.cell(td.max_row, 4).number_format = "0.0%"

    bu = wb.create_sheet("Bütçe")
    _baslik(bu, ["Kategori", "Yıllık Bütçe", f"Oransal Bütçe ({n} ay)", "Dönem Tahakkuku", "Yıl Sonu Tahmini", "Tahmin − Bütçe", "Açıklama"],
            (14, 14, 16, 15, 15, 15, 30))
    for x in s["butce"]:
        bu.append([x["kategori"], _f(x["butce"]), _f(x["oransal"]), float(x["donem"]), float(x["tahmin"]), _f(x["fark"]), ""])
        for j in range(2, 7):
            bu.cell(bu.max_row, j).number_format = PF
        if x["fark"] is not None and x["fark"] > 0:
            bu.cell(bu.max_row, 6).fill = PatternFill("solid", fgColor="FDE2E1")
        bu.cell(bu.max_row, 7).fill = PatternFill("solid", fgColor="FFF4CE")
    bu.append([])
    bu.append(["Yıl sonu tahmini = aylık kalemlerin aylık ortalaması × 12 + yıllık ödemeler + tek seferlik harcamalar (gelecek alımlar dahil değildir)."])

    li = wb.create_sheet("Lisans Kullanımı")
    _baslik(li, ["Ürün", "Toplam", "Kullanılan", "Atıl", "Kullanım", "Yıllık Birim Fiyat", "Döviz", "Atıl Lisans Maliyeti", "Yenileme", "Karar"],
            (30, 8, 10, 7, 9, 13, 6, 15, 11, 18))
    for x in s["lisans"]:
        li.append([x["urun"], float(x["toplam"]), float(x["kullanilan"]), float(x["atil"]), float(x["oran"]), _f(x["fiyat"]), x["doviz"], _f(x["atil_maliyet"]),
                   x["yenileme"], ""])
        li.cell(li.max_row, 5).number_format = "0%"
        li.cell(li.max_row, 6).number_format = li.cell(li.max_row, 8).number_format = PF
        li.cell(li.max_row, 9).number_format = "DD.MM.YYYY"
        if x["oran"] < Decimal("0.8"):
            li.cell(li.max_row, 5).fill = PatternFill("solid", fgColor="FFF4CE")
        li.cell(li.max_row, 10).fill = PatternFill("solid", fgColor="FFF4CE")

    yt = wb.create_sheet("Yenileme Takvimi")
    _baslik(yt, ["Tarih", "Kalan Gün", "Kalem / Ürün", "Tedarikçi", "Son Tutar (TL)", "Kaynak", "Aksiyon"], (11, 9, 44, 24, 14, 10, 24))
    for y in s["yenileme"]:
        yt.append([y["tarih"], (y["tarih"] - s["bugun"]).days, y["ne"], y["tedarikci"], _f(y["tutar"]), y["kaynak"], ""])
        yt.cell(yt.max_row, 1).number_format = "DD.MM.YYYY"
        yt.cell(yt.max_row, 5).number_format = PF
        yt.cell(yt.max_row, 7).fill = PatternFill("solid", fgColor="FFF4CE")

    hs = wb.create_sheet("Harcamalar")
    _baslik(hs, ["Tarih", "Tedarikçi", "Kalem", "Kategori", "Birim", "Tutar", "Döviz", "Kur", "TL", "Dönem Tipi", "Sözleşme Bitiş", "Not"],
            (11, 24, 44, 10, 13, 12, 6, 8, 13, 11, 12, 30))
    for h in sorted(s["harcamalar"], key=lambda h: (h.tarih, h.satir)):
        hs.append([h.tarih, h.tedarikci, h.kalem, h.kategori, h.birim, float(h.tutar), h.doviz, _f(h.kur), float(h.tl),
                   {"aylik": "Aylık", "yillik": "Yıllık", "tek": "Tek seferlik"}[h.tip], h.bitis, "; ".join(h.notlar)])
        hs.cell(hs.max_row, 1).number_format = hs.cell(hs.max_row, 11).number_format = "DD.MM.YYYY"
        hs.cell(hs.max_row, 6).number_format = hs.cell(hs.max_row, 9).number_format = PF
    hs.auto_filter.ref = f"A1:L{hs.max_row}"

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Kalem / Kim", "Açıklama"], (9, 24, 30, 100))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["kim"], u["aciklama"]])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(harcama_yolu: Path, cikti: Path, yil: int, bugun: date, lisans_yolu: Path | None = None, butce_yolu: Path | None = None,
             birim_yolu: Path | None = None, yenileme_gun: int = 120, artis_esik: Decimal = Decimal("0.30")) -> dict:
    harcamalar, uy = oku(harcama_yolu)
    if not harcamalar:
        raise ValueError(f"{harcama_yolu.name}: harcama bulunamadı")
    lisanslar = []
    if lisans_yolu:
        for r in kayitlar(lisans_yolu, LISANS_SUTUNLARI, ("urun", "toplam")):
            if metin(r.get("urun")):
                lisanslar.append({"urun": metin(r["urun"]), "toplam": para(r.get("toplam")) or SIFIR, "kullanilan": para(r.get("kullanilan")) or SIFIR,
                                  "fiyat": para(r.get("fiyat")), "doviz": (metin(r.get("doviz")) or "TRY").upper(), "yenileme": tarih(r.get("yenileme"))})
    butce = {}
    for r in (tablo_oku(butce_yolu)[1:] if butce_yolu else []):
        if len(r) >= 2 and para(r[1]) is not None:
            k = kategori_bul("", metin(r[0]))
            butce[k] = para(r[1])
    birimler = {metin(r[0]): int(para(r[1]) or 0) for r in (tablo_oku(birim_yolu)[1:] if birim_yolu else []) if len(r) >= 2 and metin(r[0])}
    s = analiz_et(harcamalar, yil, bugun, lisanslar, butce, birimler, yenileme_gun, artis_esik)
    s["uyarilar"] = uy + s["uyarilar"]
    rapor_yaz(cikti, s)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Lisans, bulut, donanım, hizmet ve iletişim harcamalarını birim ve kalem bazında raporlar.")
    p.add_argument("--harcamalar", type=Path, default=ORNEK / "harcamalar.csv",
                   help="Tarih, Tedarikçi, Kalem, Kategori, Birim, Tutar, Döviz, Kur, Dönem Tipi (Aylık/Yıllık/Tek seferlik), Sözleşme Bitiş")
    p.add_argument("--lisanslar", type=Path, help="İsteğe bağlı: Ürün, Toplam Lisans, Kullanılan Lisans, Yıllık Birim Fiyat, Döviz, Yenileme Tarihi")
    p.add_argument("--butce", type=Path, help="İsteğe bağlı: Kategori, Yıllık Bütçe")
    p.add_argument("--birimler", type=Path, help="İsteğe bağlı: Birim, Çalışan (kişi başı maliyet)")
    p.add_argument("--yil", type=int, help="Rapor yılı (varsayılan harcamalardaki son yıl)")
    p.add_argument("--bugun", help="Rapor tarihi GG.AA.YYYY (yenileme takvimi; örnek veride 09.10.2026)")
    p.add_argument("--yenileme-gun", type=int, default=120, help="Kaç gün içindeki yenilemeler listelensin (varsayılan 120)")
    p.add_argument("--artis", type=float, default=30, help="Aylık harcama, önceki 3 ay ortalamasını bu %%'den fazla aşarsa uyarı (varsayılan 30)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "bt_maliyet_raporu.xlsx")
    a = p.parse_args(argv)
    ornek = a.harcamalar == ORNEK / "harcamalar.csv"
    ek = {k: (getattr(a, k) or (ORNEK / f"{k}.csv" if ornek else None)) for k in ("lisanslar", "butce", "birimler")}
    for y in (a.harcamalar, *ek.values()):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    bugun = tarih(a.bugun) if a.bugun else (date(2026, 10, 9) if ornek else date.today())
    if not bugun:
        print("[X] --bugun GG.AA.YYYY biçiminde olmalı")
        return 2
    try:
        if a.yil:
            yil = a.yil
        else:
            yil = max(h.tarih.year for h in oku(a.harcamalar)[0])
        s = calistir(a.harcamalar, a.cikti, yil, bugun, ek["lisanslar"], ek["butce"], ek["birimler"], a.yenileme_gun, Decimal(str(a.artis)) / 100)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    toplam = sum((sum(v.values(), SIFIR) for v in s["kat_ay"].values()), SIFIR)
    print(f"[OK] {yil} Ocak – {AYLAR[s['son_ay'] - 1]}: tahakkuk {tl(toplam)} TL · ödenen {tl(s['nakit'])} TL · " +
          " · ".join(f"{k} {tl(sum(v.values(), SIFIR))}" for k, v in sorted(s["kat_ay"].items(), key=lambda i: -sum(i[1].values(), SIFIR))))
    for u in s["uyarilar"]:
        if u["onem"] != "Bilgi":
            print(f"[!] {u['kim']} · {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
