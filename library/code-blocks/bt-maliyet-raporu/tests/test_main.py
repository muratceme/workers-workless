import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "b.xlsx"
        cls.s = main.calistir(ORNEK / "harcamalar.csv", cls.cikti, 2026, date(2026, 10, 9), ORNEK / "lisanslar.csv", ORNEK / "butce.csv", ORNEK / "birimler.csv")
        cls.u = {(u["tur"], u["kim"]) for u in cls.s["uyarilar"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_kategoriler(self):
        k = {h.kalem: h.kategori for h in self.s["harcamalar"]}
        self.assertEqual(k["Bulut sunucu ve depolama (aylık kullanım)"], "Bulut")
        self.assertEqual(k["ERP lisans bakım bedeli"], "Lisans")
        self.assertEqual(k["Sunucu disk ve bellek yükseltme"], "Donanım")
        self.assertEqual(k["Mobil hat paketi"], "İletişim")
        self.assertEqual(k["Sızma testi hizmeti"], "Hizmet")

    def test_tahakkuk_ve_korunum(self):
        tahakkuk = sum(sum(v.values()) for v in self.s["kat_ay"].values())
        self.assertAlmostEqual(float(tahakkuk + sum(self.s["pesin"].values())), float(self.s["nakit"]), places=2)
        erp = next(h for h in self.s["harcamalar"] if h.kalem.startswith("ERP"))
        d, kalan = main.tahakkuk(erp, 2026)
        self.assertEqual((len(d), kalan), (12, D("0")))                  # Ocak'tan başlayan yıllık: 12 ay rapor yılı içinde
        cad = next(h for h in self.s["harcamalar"] if h.kalem.startswith("CAD"))
        d, kalan = main.tahakkuk(cad, 2026)
        self.assertEqual((min(d), len(d)), (3, 10))                      # Mart–Aralık; 2 ay gelecek yıla
        self.assertAlmostEqual(float(kalan), float(cad.tl / 6), places=2)

    def test_uyarilar(self):
        self.assertTrue({("Harcama artışı", "Örnek Bulut Hizmetleri · Bilgi İşlem"), ("Olası mükerrer abonelik", "Örnek Tasarım SaaS"),
                         ("Atıl lisans", "CAD yazılımı"), ("Kur tahmini", "Güvenlik duvarı lisans yenileme"), ("Bütçe aşım tahmini", "Lisans")} <= self.u)
        self.assertNotIn(("Atıl lisans", "Microsoft 365 Business Standard"), self.u)    # tam %80
        self.assertEqual(len(self.s["yenileme"]), 4)

    def test_excel(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Özet", "Kategori × Ay", "Birim × Kategori", "Tedarikçiler", "Bütçe", "Lisans Kullanımı", "Yenileme Takvimi",
                                         "Harcamalar", "Uyarılar"])
        self.assertEqual([r[0] for r in wb["Kategori × Ay"].iter_rows(min_row=2, values_only=True)][-1], "Toplam")
        self.assertNotIn("Diğer", [r[0] for r in wb["Kategori × Ay"].iter_rows(min_row=2, values_only=True)])


class KuralTesti(unittest.TestCase):
    def test_kategori_bul(self):
        self.assertEqual(main.kategori_bul("Azure aboneliği"), "Bulut")
        self.assertEqual(main.kategori_bul("Kalem", "Donanım"), "Donanım")
        self.assertEqual(main.kategori_bul("Ofis kahvesi"), "Diğer")

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--bugun", "x"]), 2)
            self.assertEqual(main.main(["--harcamalar", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
