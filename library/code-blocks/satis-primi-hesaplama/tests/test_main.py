import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
DILIM = [(D(0), D(0)), (D(80), D(1)), (D(100), D("1.5")), (D(120), D(2))]


def k(satis, hedef=1000, tahsilat=None, ad="A"):
    d = lambda x: None if x is None else D(str(x))   # noqa: E731
    return {"ad": ad, "hedef": d(hedef), "satis": d(satis), "tahsilat": d(tahsilat)}


def h(kisi, dilimler=DILIM, **kw):
    r = main.hesapla([kisi], dilimler, **kw)
    x = r["kisiler"][0]
    return x["oran_g"], x["kademe"], x["prim_orani"], x["ham"], x["prim"], [(u["onem"], u["tur"]) for u in r["uyarilar"]]


class HesapTesti(unittest.TestCase):
    def test_kademe(self):
        self.assertEqual([h(k(s))[:5] for s in (800, "799.99", 1000, "1199.99", 1200, 0)],
                         [(D("80.00"), "%80 – %100", D(1), D("8.00"), D("8.00")), (D("80.00"), "%0 – %80", D(0), D("0.00"), D("0.00")), (D("100.00"), "%100 – %120", D("1.5"), D("15.00"), D("15.00")),
                          (D("120.00"), "%100 – %120", D("1.5"), D("18.00"), D("18.00")), (D("120.00"), "%120 ve üstü", D(2), D("24.00"), D("24.00")), (D("0.00"), "%0 – %80", D(0), D("0.00"), D("0.00"))])
        self.assertEqual(h(k(700), [(D(80), D(1)), (D(100), D(2))]), (D("70.00"), None, None, D(0), D(0), []))
        self.assertEqual([h(k(100, x)) for x in (None, 0)], [(None, None, None, D(0), D(0), [("Yüksek", "Hedef yok")])] * 2)
        r = main.hesapla([k(1250)], DILIM)
        self.assertEqual((r["kisiler"][0]["fark"], r["dokum"]), (D(250), [{"ad": "A", "dilim": "%120 ve üstü", "tutar": D(1250), "oran": D(2), "prim": D("25.00")}]))

    def test_dilimli(self):
        r = main.hesapla([k(1250)], DILIM, "dilimli")
        self.assertEqual([(d["dilim"], d["tutar"], d["oran"], d["prim"]) for d in r["dokum"]],
                         [("%0 – %80", D("800.00"), D(0), D("0.00")), ("%80 – %100", D("200.00"), D(1), D("2.00")), ("%100 – %120", D("200.00"), D("1.5"), D("3.00")), ("%120 ve üstü", D("50.00"), D(2), D("1.00"))])
        x = r["kisiler"][0]
        self.assertEqual((x["ham"], x["prim"], x["prim_orani"], x["kademe"]), (D("6.00"), D("6.00"), D("0.48"), "%120 ve üstü"))
        self.assertEqual([h(k(s), yontem="dilimli")[2:5] for s in (900, 800, 700, 0)], [(D("0.11"), D("1.00"), D("1.00")), (D("0.00"), D("0.00"), D("0.00")), (D("0.00"), D("0.00"), D("0.00")), (None, D(0), D(0))])
        self.assertEqual(h(k(100000, 300000), [(D(0), D(3)), (D(30), D(6))], yontem="dilimli")[3], D("3300.00"))
        self.assertEqual((h(k(950), [(D(90), D(10))], yontem="dilimli")[3], h(k(850), [(D(90), D(10))], yontem="dilimli")[3], main.hesapla([k(850)], [(D(90), D(10))], "dilimli")["dokum"]), (D("5.00"), D(0), []))

    def test_tahsilat_ve_tavan(self):
        t = lambda tahsilat, **kw: (lambda r: (r["kisiler"][0]["tahsilat_orani"], r["kisiler"][0]["ham"], r["kisiler"][0]["prim"], r["kisiler"][0]["nota"], [(u["onem"], u["tur"]) for u in r["uyarilar"]]))(   # noqa: E731
            main.hesapla([k(1250, tahsilat=tahsilat)], DILIM, **kw))
        self.assertEqual(t(1100, esas="tahsilat"), (D("88.00"), D("25.00"), D("22.00"), "", []))
        self.assertEqual(t(1250, esas="tahsilat")[:3], (D("100.00"), D("25.00"), D("25.00")))
        self.assertEqual(t(1300, esas="tahsilat"), (D("100.00"), D("25.00"), D("25.00"), "", [("Bilgi", "Tahsilat satıştan büyük")]))
        self.assertEqual(t(None, esas="tahsilat"), (D("0.00"), D("25.00"), D("0.00"), "Tahsilat yazılmamış", [("Orta", "Tahsilat yazılmamış")]))
        self.assertEqual(t(1100), (None, D("25.00"), D("25.00"), "", []))
        self.assertEqual(t(1100, tavan=D(20)), (None, D("25.00"), D(20), "Tavan uygulandı (25.00 → 20.00)", [("Bilgi", "Prim tavanı uygulandı")]))
        self.assertEqual((t(1100, tavan=D(25))[2:], t(1100, esas="tahsilat", tavan=D(22))[2:], t(1100, esas="tahsilat", tavan=D(21))[2]), ((D("25.00"), "", []), (D("22.00"), "", []), D(21)))
        self.assertEqual(main.hesapla([k(0, tahsilat=0)], DILIM, esas="tahsilat")["kisiler"][0]["tahsilat_orani"], D("0.00"))
        with self.assertRaises(ValueError):
            main.hesapla([], DILIM)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "gerceklesme.csv", ORNEK / "dilimler.csv", cikti)
            self.assertEqual(r["dilimler"], DILIM)
            self.assertEqual([(x["ad"], x["oran_g"], x["fark"], x["kademe"], x["prim_orani"], x["prim"]) for x in r["kisiler"]],
                             [("Burak Şen", D("125.00"), D(250000), "%120 ve üstü", D(2), D(25000)), ("Deniz Acar", D("95.00"), D(-40000), "%80 – %100", D(1), D(7600)), ("Ece Yurt", D("70.00"), D(-180000), "%0 – %80", D(0), D(0)),
                              ("Can Polat", D("100.00"), D(0), "%100 – %120", D("1.5"), D(13500)), ("Selin Arı", D("142.00"), D(210000), "%120 ve üstü", D(2), D(14200)), ("Okan Tuna", None, None, None, None, D(0))])
            self.assertEqual(r["toplam"], {"hedef": D(3800000), "satis": D(4190000), "oran_g": D("106.32"), "prim": D(60300), "prim_satis": D("1.44"), "prim_alan": 4})
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]], [("Yüksek", "Hedef yok", "Okan Tuna")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Primler", "Dilim Dökümü", "Prim Sistemi", "Uyarılar"])
            r = main.calistir(ORNEK / "gerceklesme.csv", ORNEK / "dilimler.csv", cikti, "dilimli")
            self.assertEqual(([x["prim"] for x in r["kisiler"]], r["toplam"]["prim"], r["toplam"]["prim_satis"], r["toplam"]["prim_alan"], len(r["dokum"])), ([D(6000), D(1200), D(0), D(1800), D(4700), D(0)], D(13700), D("0.33"), 4, 13))
            r = main.calistir(ORNEK / "gerceklesme.csv", ORNEK / "dilimler.csv", cikti, "kademe", "tahsilat", D(20000))
            self.assertEqual(([(x["tahsilat_orani"], x["prim"]) for x in r["kisiler"][:5]], r["toplam"]["prim"], sorted(u["tur"] for u in r["uyarilar"])),
                             ([(D("88.00"), D(20000)), (D("100.00"), D(7600)), (D("95.24"), D(0)), (D("70.00"), D(9450)), (D("100.00"), D(14200))], D(51250), ["Hedef yok", "Prim tavanı uygulandı"]))

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            ky, dy = Path(tmp) / "k.csv", Path(tmp) / "d.csv"
            ky.write_text("Personel;Kota;Satış;Tahsil Edilen\nAli;1.000.000;1.250.000,50;\nEce;;5;\nCan;x;5;\nOya;5;-1;\nEda;5;5;x\n;5;5;5\n", encoding="utf-8")
            dy.write_text("Gerçekleşme;Prim\n100;1,5\n80;1\n;\n%120;2.5\n", encoding="utf-8")
            l, uy = main.oku_kisiler(ky)
            self.assertEqual((l, [(u["tur"], u["kim"]) for u in uy]),
                             ([{"ad": "Ali", "hedef": D(1000000), "satis": D("1250000.50"), "tahsilat": None}, {"ad": "Ece", "hedef": None, "satis": D(5), "tahsilat": None}], [("Satır okunamadı", x) for x in ("Can", "Oya", "Eda")]))
            self.assertEqual(main.oku_dilimler(dy), [(D(80), D(1)), (D(100), D("1.5")), (D(120), D("2.5"))])
            self.assertEqual([main.olcu(x) for x in ("1,5", "1.5", "%2", "1.250,5", "", "x", 3)], [D("1.5"), D("1.5"), D(2), D("1250.5"), None, None, D(3)])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--gerceklesme", str(ky), "--dilimler", str(dy), "--cikti", c]
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(arg + ["--yontem", "dilimli", "--esas", "tahsilat", "--tavan", "20.000"]), 0)
            self.assertEqual([main.main(arg + x) for x in (["--tavan", "0"], ["--tavan", "x"])] + [main.main(["--gerceklesme", str(ky)])], [2] * 3)
            self.assertEqual(main.main(["--gerceklesme", str(ky), "--dilimler", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Alt Sınır;Prim Oranı\n80;1\n80;2\n", "Alt Sınır;Prim Oranı\n80;-1\n", "Alt Sınır;Prim Oranı\nx;1\n", "Alt Sınır;Prim Oranı\n;\n", "Alt Sınır\n80\n"):
                dy.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            dy.write_text("Alt Sınır;Prim Oranı\n80;1\n", encoding="utf-8")
            for icerik in ("Temsilci;Hedef\nA;1\n", "Temsilci;Gerçekleşen\n;1\n"):
                ky.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
