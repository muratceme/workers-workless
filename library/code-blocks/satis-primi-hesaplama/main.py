"""
Satış Primi Hesaplama — Workers / Workless kod bloğu
Satış › Satış Müdürü

Prim sistemine göre temsilci bazında hedef gerçekleşmeyi ve brüt prim tutarını hesaplar.
  Gerçekleşme %      = gerçekleşen satış ÷ hedef × 100
  Kademe yöntemi     = gerçekleşme yüzdesinin düştüğü dilimin oranı satışın tamamına uygulanır
                       prim = gerçekleşen satış × dilim oranı
  Dilimli yöntem     = her dilime düşen satış kendi oranıyla primlendirilir (artan oranlı)
                       dilime düşen satış = min(gerçekleşen, hedef × dilim üst sınırı) − hedef × dilim alt sınırı
  Tahsilat esası     = prim × min(tahsilat ÷ gerçekleşen satış, 1)   (--esas tahsilat)
  Tavan              = prim, --tavan tutarını aşamaz (verilirse)
Dilimler ve oranlar sizin prim sisteminizdir; dosyadan okunur. Sonuç brüt primdir: bordro kesintileri hesaplanmaz.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
K_SUTUN = {"ad": ("temsilci", "satis temsilcisi", "personel", "ad soyad"), "hedef": ("hedef", "satis hedefi", "kota"), "satis": ("gerceklesen", "gerceklesen satis", "satis", "ciro"), "tahsilat": ("tahsilat", "tahsil edilen")}
D_SUTUN = {"alt": ("alt sinir", "alt sinir %", "gerceklesme alt siniri", "gerceklesme", "baslangic"), "oran": ("prim orani", "prim orani %", "oran", "prim")}
Z, YUZ = Decimal(0), Decimal(100)


def olcu(x) -> Decimal | None:
    """Yüzde ve oran okur: virgül ya da tek nokta ondalık ayırıcıdır (1,5 = 1.5)."""
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x or ""))
    try:
        return Decimal(s.replace(".", "").replace(",", ".") if "," in s else s) if s else None
    except InvalidOperation:
        return None


def oku_kisiler(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, K_SUTUN, ("ad", "satis")):
        ad = metin(r["ad"])
        if not ad:
            continue
        h, s, t = para(r["hedef"]), para(r["satis"]), para(r["tahsilat"])
        if s is None or s < 0 or (h is not None and h < 0) or (t is not None and t < 0) or any(metin(r[a]) and v is None for a, v in (("hedef", h), ("tahsilat", t))):
            uy.append(uyari("Yüksek", "Satır okunamadı", ad, f"Satır {r['_satir']}: hedef '{metin(r['hedef'])}', gerçekleşen '{metin(r['satis'])}', tahsilat '{metin(r['tahsilat'])}'. Eksi ya da okunamayan değer var; satır alınmadı."))
            continue
        l.append({"ad": ad, "hedef": h, "satis": s, "tahsilat": t})
    return l, uy


def oku_dilimler(yol: Path) -> list[tuple[Decimal, Decimal]]:
    d = []
    for r in kayitlar(yol, D_SUTUN, ("alt", "oran")):
        if not metin(r["alt"]) and not metin(r["oran"]):
            continue
        alt, oran = olcu(r["alt"]), olcu(r["oran"])
        if alt is None or oran is None or alt < 0 or oran < 0:
            raise ValueError(f"{yol.name} satır {r['_satir']}: alt sınır ve prim oranı sıfır veya pozitif sayı olmalı.")
        d.append((alt, oran))
    d.sort()
    if not d or any(a[0] == b[0] for a, b in zip(d, d[1:])):
        raise ValueError(f"{yol.name}: en az bir dilim olmalı ve aynı alt sınır iki kez yazılmamalı.")
    return d


def dilim_adi(dilimler: list, i: int) -> str:
    return f"%{dilimler[i][0]:f} – %{dilimler[i + 1][0]:f}" if i + 1 < len(dilimler) else f"%{dilimler[i][0]:f} ve üstü"


def hesapla(kisiler: list[dict], dilimler: list[tuple[Decimal, Decimal]], yontem: str = "kademe", esas: str = "satis", tavan: Decimal | None = None) -> dict:
    if not kisiler:
        raise ValueError("Prim hesaplanacak temsilci yok.")
    uy, dokum = [], []
    for k in kisiler:
        k["oran_g"] = k["fark"] = k["kademe"] = k["prim_orani"] = None
        k["ham"], k["tahsilat_orani"], k["prim"], k["nota"] = Z, None, Z, ""
        if not k["hedef"]:
            k["nota"] = "Hedef yok"
            uy.append(uyari("Yüksek", "Hedef yok", k["ad"], f"Hedef yazılmamış ya da sıfır; {k['satis']:,.2f} satış için prim hesaplanamadı."))
            continue
        g = k["satis"] * YUZ / k["hedef"]
        k["oran_g"], k["fark"] = yuvarla(g), k["satis"] - k["hedef"]
        i = max((j for j, d in enumerate(dilimler) if g >= d[0]), default=None)
        if i is not None:
            k["kademe"] = dilim_adi(dilimler, i)
        if yontem == "kademe":
            if i is not None:
                k["prim_orani"] = dilimler[i][1]
                k["ham"] = yuvarla(k["satis"] * dilimler[i][1] / YUZ)
                dokum.append({"ad": k["ad"], "dilim": k["kademe"], "tutar": k["satis"], "oran": dilimler[i][1], "prim": k["ham"]})
        else:
            for j, (alt, oran) in enumerate(dilimler):
                ust = min(k["satis"], k["hedef"] * dilimler[j + 1][0] / YUZ) if j + 1 < len(dilimler) else k["satis"]
                pay = ust - k["hedef"] * alt / YUZ
                if pay > 0:
                    dokum.append({"ad": k["ad"], "dilim": dilim_adi(dilimler, j), "tutar": yuvarla(pay), "oran": oran, "prim": yuvarla(pay * oran / YUZ)})
                    k["ham"] += dokum[-1]["prim"]
            k["prim_orani"] = yuvarla(k["ham"] * YUZ / k["satis"]) if k["satis"] else None
        k["prim"], notlar = k["ham"], []
        if esas == "tahsilat":
            if k["tahsilat"] is None:
                k["tahsilat_orani"] = Z
                notlar.append("Tahsilat yazılmamış")
                uy.append(uyari("Orta", "Tahsilat yazılmamış", k["ad"], "Prim tahsilat esaslı hesaplanıyor ama tahsilat sütunu boş; prim sıfır alındı."))
            else:
                k["tahsilat_orani"] = min(k["tahsilat"] / k["satis"], Decimal(1)) if k["satis"] else Z
                if k["tahsilat"] > k["satis"]:
                    uy.append(uyari("Bilgi", "Tahsilat satıştan büyük", k["ad"], f"Tahsilat {k['tahsilat']:,.2f}, satış {k['satis']:,.2f}. Önceki dönem tahsilatı karışmış olabilir; oran %100 ile sınırlandı."))
            k["prim"] = yuvarla(k["ham"] * k["tahsilat_orani"])
            k["tahsilat_orani"] = yuvarla(k["tahsilat_orani"] * YUZ)
        if tavan is not None and k["prim"] > tavan:
            notlar.append(f"Tavan uygulandı ({k['prim']:,.2f} → {tavan:,.2f})")
            uy.append(uyari("Bilgi", "Prim tavanı uygulandı", k["ad"], f"Hesaplanan prim {k['prim']:,.2f}, tavan {tavan:,.2f}."))
            k["prim"] = tavan
        k["nota"] = "; ".join(notlar)
    hedefli = [k for k in kisiler if k["hedef"]]
    hedef, satis, hedefli_satis, prim = sum((k["hedef"] for k in hedefli), Z), sum((k["satis"] for k in kisiler), Z), sum((k["satis"] for k in hedefli), Z), sum((k["prim"] for k in kisiler), Z)
    return {"kisiler": kisiler, "dokum": dokum, "dilimler": dilimler, "yontem": yontem, "esas": esas, "tavan": tavan,
            "toplam": {"hedef": hedef, "satis": satis, "oran_g": yuvarla(hedefli_satis * YUZ / hedef) if hedef else None, "prim": prim, "prim_satis": yuvarla(prim * YUZ / satis) if satis else None,
                       "prim_alan": sum(1 for k in kisiler if k["prim"] > 0)}, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Primler"
    _baslik(ws, ["Temsilci", "Hedef", "Gerçekleşen Satış", "Gerçekleşme %", "Hedef Farkı", "Kademe", "Prim Oranı %", "Hesaplanan Prim", "Tahsilat", "Tahsilat Oranı %", "Brüt Prim", "Not"], (26, 16, 16, 11, 16, 20, 10, 15, 16, 11, 15, 40))
    for k in r["kisiler"]:
        _satir(ws, [k["ad"], k["hedef"], k["satis"], k["oran_g"], k["fark"], k["kademe"], k["prim_orani"], k["ham"], k["tahsilat"], k["tahsilat_orani"], k["prim"], k["nota"]], para_s=(2, 3, 4, 5, 7, 8, 9, 10, 11),
               dolgu={4: None if k["oran_g"] is None else YESIL if k["oran_g"] >= 100 else SARI, 11: YESIL if k["prim"] > 0 else None})
    t = r["toplam"]
    n = _satir(ws, ["TOPLAM", t["hedef"], t["satis"], t["oran_g"], None, None, t["prim_satis"], None, None, None, t["prim"], f"{t['prim_alan']} kişi prim alıyor"], para_s=(2, 3, 4, 7, 11))
    for h in ws[n]:
        h.font = Font(bold=True)
    ws.append([])
    ws.append([f"Yöntem: {'kademe (dilim oranı satışın tamamına)' if r['yontem'] == 'kademe' else 'dilimli (her dilime düşen satış kendi oranıyla)'} · esas: {'satış' if r['esas'] == 'satis' else 'tahsilat'}"
               + (f" · kişi başı tavan {r['tavan']:,.2f}" if r["tavan"] is not None else "") + " · Tutarlar brüttür."])
    ws = wb.create_sheet("Dilim Dökümü")
    _baslik(ws, ["Temsilci", "Dilim (gerçekleşme)", "Dilime Düşen Satış", "Prim Oranı %", "Prim"], (26, 22, 18, 11, 15))
    for d in r["dokum"]:
        _satir(ws, [d["ad"], d["dilim"], d["tutar"], d["oran"], d["prim"]], para_s=(3, 4, 5))
    ws = wb.create_sheet("Prim Sistemi")
    _baslik(ws, ["Gerçekleşme Dilimi", "Prim Oranı %"], (24, 12))
    for i, d in enumerate(r["dilimler"]):
        _satir(ws, [dilim_adi(r["dilimler"], i), d[1]], para_s=(2,))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(k_yol: Path, d_yol: Path, cikti: Path, yontem: str = "kademe", esas: str = "satis", tavan: Decimal | None = None) -> dict:
    kisiler, uy = oku_kisiler(k_yol)
    r = hesapla(kisiler, oku_dilimler(d_yol), yontem, esas, tavan)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Prim sistemine göre temsilci bazında hedef gerçekleşmeyi ve brüt primi hesaplar.")
    p.add_argument("--gerceklesme", type=Path, help="Temsilci, Hedef, Gerçekleşen, Tahsilat")
    p.add_argument("--dilimler", type=Path, help="Prim sistemi: Alt Sınır (gerçekleşme yüzdesi), Prim Oranı (yüzde)")
    p.add_argument("--yontem", choices=("kademe", "dilimli"), default="kademe", help="kademe: dilim oranı satışın tamamına (varsayılan); dilimli: her dilime düşen satış kendi oranıyla")
    p.add_argument("--esas", choices=("satis", "tahsilat"), default="satis", help="tahsilat: prim, satışın tahsil edilen kısmı oranında ödenir")
    p.add_argument("--tavan", help="Kişi başı prim tavanı (tutar); verilmezse tavan yok")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "satis_primi.xlsx")
    a = p.parse_args(argv)
    if (a.gerceklesme is None) != (a.dilimler is None):
        print("[X] --gerceklesme ve --dilimler birlikte verilmelidir")
        return 2
    k_yol, d_yol = (ORNEK / "gerceklesme.csv", ORNEK / "dilimler.csv") if a.gerceklesme is None else (a.gerceklesme, a.dilimler)
    tavan = para(a.tavan)
    if a.tavan and (tavan is None or tavan <= 0):
        print("[X] --tavan pozitif bir tutar olmalı")
        return 2
    for y in (k_yol, d_yol):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(k_yol, d_yol, a.cikti, a.yontem, a.esas, tavan)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {len(r['kisiler'])} temsilci · satış {t['satis']:,.2f}" + (f" · gerçekleşme %{t['oran_g']}" if t["oran_g"] is not None else "") + f" · brüt prim {t['prim']:,.2f}" + (f" (satışın %{t['prim_satis']}'i)" if t["prim_satis"] is not None else "")
          + f" · {t['prim_alan']} kişi prim alıyor")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
