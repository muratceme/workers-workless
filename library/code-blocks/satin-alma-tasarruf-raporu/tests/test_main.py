import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def a(ad, miktar, butce, gercek, teklif=None, kategori="K", tur="Pazarlık", kisi="A", ay=7):
    d = lambda x: None if x is None else D(str(x))   # noqa: E731
    return {"tarih": date(2026, ay, 10), "ay": f"2026-{ay:02d}", "kategori": kategori, "ad": ad, "tedarikci": "T", "miktar": D(miktar), "butce": d(butce), "teklif": d(teklif), "gercek": d(gercek), "tur": tur, "kisi": kisi}


class RaporTesti(unittest.TestCase):
    def test_satir_hesabi(self):
        r = main.raporla([a("tasarruf", 1000, 75, 64, 72), a("aşım", 10, 100, 110), a("bütçesiz", 4, None, 9000, 9500), a("pahalı", 2, 50, 50, 45), a("ondalık", 3, "0.335", "0.33")])
        self.assertEqual([(x["butce_tutar"], x["gercek_tutar"], x["tasarruf"], x["oran"], x["pazarlik"]) for x in r["alimlar"]],
                         [(D(75000), D(64000), D(11000), D("14.67"), D(8000)), (D(1000), D(1100), D(-100), D("-10.00"), None), (None, D(36000), None, None, D(2000)), (D(100), D(100), D(0), D("0.00"), D(-10)),
                          (D("1.01"), D("0.99"), D("0.02"), D("1.98"), None)])
        self.assertEqual((r["toplam"], r["butcesiz"]),
                         ({"kalem": 5, "harcama": D("101200.99"), "butce": D("76101.01"), "butceli_harcama": D("65200.99"), "brut": D("11000.02"), "asim": D(100), "pazarlik": D(9990), "net": D("10900.02"), "oran": D("14.32")}, 1))
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]), [("Bütçe aşımı", "aşım · T"), ("Bütçe fiyatı yok", "bütçesiz · T"), ("İlk tekliften pahalı alım", "pahalı · T")])

    def test_esikler_ve_kirilim(self):
        u = lambda l, **kw: [x["tur"] for x in main.raporla(l, **kw)["uyarilar"]]   # noqa: E731
        self.assertEqual([u([a("x", 100, 100, g)]) for g in (105, 104.99, 60, 60.01)], [["Bütçe aşımı"], [], ["Tasarruf oranı çok yüksek"], []])
        self.assertEqual((u([a("x", 100, 100, 103)], asim_esik=D(3)), u([a("x", 100, 100, 70)], supheli=D(30)), u([a("x", 100, 100, 100)], asim_esik=D(0))), (["Bütçe aşımı"], ["Tasarruf oranı çok yüksek"], []))
        l = [a("1", 10, 100, 90, kategori="Ambalaj", tur="Pazarlık", kisi="Ali", ay=7), a("2", 10, 100, 120, kategori="ambalaj", tur="", kisi="Ece", ay=8), a("3", 10, 50, 40, 45, kategori="Hammadde", tur="Alternatif", kisi="Ali", ay=8),
             a("4", 1, None, 500, kategori="Hammadde", tur="Pazarlık", kisi="Ece", ay=7)]
        r = main.raporla(l)
        k = lambda ad: [(s["ad"], s["kalem"], s["harcama"], s["butce"], s["brut"], s["asim"], s["net"], s["oran"]) for s in r["kirilimlar"][ad]]   # noqa: E731
        self.assertEqual(k("Kategori"), [("Hammadde", 2, D(900), D(500), D(100), D(0), D(100), D("20.00")), ("Ambalaj", 2, D(2100), D(2000), D(100), D(200), D(-100), D("-5.00"))])
        self.assertEqual(k("Tasarruf Türü"), [("Alternatif", 1, D(400), D(500), D(100), D(0), D(100), D("20.00")), ("Pazarlık", 2, D(1400), D(1000), D(100), D(0), D(100), D("10.00")), ("", 1, D(1200), D(1000), D(0), D(200), D(-200), D("-20.00"))])
        self.assertEqual(k("Satın Almacı"), [("Ali", 2, D(1300), D(1500), D(200), D(0), D(200), D("13.33")), ("Ece", 2, D(1700), D(1000), D(0), D(200), D(-200), D("-20.00"))])
        self.assertEqual(k("Ay"), [("2026-07", 2, D(1400), D(1000), D(100), D(0), D(100), D("10.00")), ("2026-08", 2, D(1600), D(1500), D(100), D(200), D(-100), D("-6.67"))])
        self.assertEqual(r["kirilimlar"]["Kategori"][0]["pazarlik"], D(50))
        self.assertEqual(main.raporla([a("x", 1, None, 5)])["toplam"]["oran"], None)
        with self.assertRaises(ValueError):
            main.raporla([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "alimlar.csv", cikti)
            self.assertEqual((r["toplam"], r["butcesiz"]),
                             ({"kalem": 8, "harcama": D(797300), "butce": D(801000), "butceli_harcama": D(761300), "brut": D(49300), "asim": D(9600), "pazarlik": D(76300), "net": D(39700), "oran": D("4.96")}, 1))
            self.assertEqual([(s["ad"], s["harcama"], s["butce"], s["net"], s["oran"]) for s in r["kirilimlar"]["Kategori"]],
                             [("Ambalaj", D(156000), D(179000), D(23000), D("12.85")), ("Hizmet", D(228000), D(240000), D(12000), D("5.00")), ("Hammadde", D(343100), D(346000), D(2900), D("0.84")),
                              ("Yedek Parça", D(70200), D(36000), D(1800), D("5.00"))])
            self.assertEqual([(s["ad"], s["net"]) for s in r["kirilimlar"]["Tasarruf Türü"]], [("Pazarlık", D(27500)), ("Alternatif Tedarikçi", D(20000)), ("Toplu Alım", D(1800)), ("(belirtilmemiş)", D(-9600))])
            self.assertEqual([(s["ad"], s["net"], s["oran"]) for s in r["kirilimlar"]["Satın Almacı"]], [("Zeynep Ak", D(35000), D("8.35")), ("Okan Er", D(4700), D("1.23"))])
            self.assertEqual([(s["ad"], s["net"], s["oran"]) for s in r["kirilimlar"]["Ay"]], [("2026-07", D(5900), D("1.37")), ("2026-08", D(24800), D("7.07")), ("2026-09", D(9000), D("45.00"))])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Orta", "Bütçe aşımı", "Boya Pigmenti · Kimya H Ltd."), ("Orta", "Bütçe fiyatı yok", "Konveyör Bandı · Teknik J Ltd."), ("Bilgi", "Tasarruf oranı çok yüksek", "Ürün Etiketi · Matbaa D Ltd.")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Kategori", "Tasarruf Türü", "Satın Almacı", "Ay", "Alımlar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "a.csv"
            y.write_text("Sipariş Tarihi;Grup;Ürün;Satıcı;Adet;Hedef Fiyat;Teklif Fiyatı;Alım Fiyatı;Yöntem;Alıcı\n05.07.2026;G;Koli;T;1.500;5,60;;5,40;;\nx;;Bant;T;10;0;-1;3;İhale;Ali\n05.07.2026;G;Film;T;0;5;5;5;;\n"
                         "05.07.2026;G;Etiket;T;5;5;5;x;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["ad"], x["ay"], x["kategori"], x["miktar"], x["butce"], x["teklif"], x["gercek"], x["tur"], x["kisi"]) for x in l], [u["tur"] for u in uy]),
                             ([("Koli", "2026-07", "G", D(1500), D("5.60"), None, D("5.40"), "(belirtilmemiş)", "(belirtilmemiş)"), ("Bant", "(tarih yok)", "(kategori yok)", D(10), None, None, D(3), "İhale", "Ali")],
                              ["Alım okunamadı"] * 2))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--alimlar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--asim", "3", "--supheli", "30"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + x) for x in (["--asim=-1"], ["--asim", "x"], ["--supheli", "0"], ["--supheli", "101"])], [2] * 4)
            self.assertEqual(main.main(["--alimlar", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Malzeme;Miktar\nA;1\n", "Malzeme;Miktar;Gerçekleşen Fiyat\nA;x;1\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
