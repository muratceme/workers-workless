"""
Satın Alma Tasarruf Raporu — Workers / Workless kod bloğu
Satın Alma › Satın Alma Müdürü

Gerçekleşen alımları bütçe fiyatıyla karşılaştırır; pazarlık ve alternatif tedarik kaynaklı tasarrufu raporlar.
  Bütçe tutarı        = miktar × bütçe fiyatı;   gerçekleşen tutar = miktar × gerçekleşen fiyat
  Tasarruf            = bütçe tutarı − gerçekleşen tutar   (eksi değer = bütçe aşımı)
  Tasarruf oranı %    = tasarruf ÷ bütçe tutarı × 100
  Brüt tasarruf       = artı tasarrufların toplamı;   net tasarruf = brüt tasarruf − aşımlar
  Pazarlık kazancı    = (ilk teklif fiyatı − gerçekleşen fiyat) × miktar   (ilk teklif yazılmışsa; ayrı bir ölçüdür, bütçeyle ilgisi yoktur)
  Kırılımlar          = kategori, tasarruf türü, satın almacı ve ay bazında
Bütçe fiyatı olmayan alım harcamaya girer ama tasarruf hesabına girmez. Tasarrufun gerçekliği bütçe fiyatının doğruluğuna bağlıdır.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
SUTUN = {"tarih": ("tarih", "siparis tarihi", "alim tarihi"), "kategori": ("kategori", "malzeme grubu", "grup"), "ad": ("malzeme", "malzeme adi", "kalem", "urun", "hizmet"), "tedarikci": ("tedarikci", "satici", "firma"),
         "miktar": ("miktar", "adet"), "butce": ("butce fiyati", "butce birim fiyat", "referans fiyat", "hedef fiyat"), "teklif": ("ilk teklif fiyati", "ilk teklif", "teklif fiyati"),
         "gercek": ("gerceklesen fiyat", "gerceklesen birim fiyat", "alim fiyati", "birim fiyat"), "tur": ("tasarruf turu", "tasarruf kaynagi", "yontem"), "kisi": ("satin almaci", "satinalmaci", "sorumlu", "alici")}
Z, YUZ = Decimal(0), Decimal(100)
KIRILIM = (("Kategori", "kategori"), ("Tasarruf Türü", "tur"), ("Satın Almacı", "kisi"), ("Ay", "ay"))


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("ad", "miktar", "gercek")):
        ad = metin(r["ad"])
        if not ad:
            continue
        m, g, b, tk, t = para(r["miktar"]), para(r["gercek"]), para(r["butce"]), para(r["teklif"]), tarih(r["tarih"])
        if m is None or m <= 0 or g is None or g <= 0:
            uy.append(uyari("Yüksek", "Alım okunamadı", f"{ad} (satır {r['_satir']})", f"Miktar '{metin(r['miktar'])}', gerçekleşen fiyat '{metin(r['gercek'])}'. İkisi de pozitif olmalı. Satır rapora alınmadı."))
            continue
        l.append({"tarih": t, "ay": f"{t:%Y-%m}" if t else "(tarih yok)", "kategori": metin(r["kategori"]) or "(kategori yok)", "ad": ad, "tedarikci": metin(r["tedarikci"]), "miktar": m, "butce": b if b and b > 0 else None,
                  "teklif": tk if tk and tk > 0 else None, "gercek": g, "tur": metin(r["tur"]) or "(belirtilmemiş)", "kisi": metin(r["kisi"]) or "(belirtilmemiş)"})
    return l, uy


def _oran(tasarruf: Decimal, butce: Decimal) -> Decimal | None:
    return yuvarla(tasarruf * YUZ / butce) if butce else None


def raporla(alimlar: list[dict], asim_esik: Decimal = Decimal(5), supheli: Decimal = Decimal(40)) -> dict:
    if not alimlar:
        raise ValueError("Raporlanacak alım yok.")
    uy = []
    for a in alimlar:
        a["gercek_tutar"] = yuvarla(a["miktar"] * a["gercek"])
        a["butce_tutar"] = yuvarla(a["miktar"] * a["butce"]) if a["butce"] else None
        a["tasarruf"] = a["butce_tutar"] - a["gercek_tutar"] if a["butce"] else None
        a["oran"] = _oran(a["tasarruf"], a["butce_tutar"]) if a["butce"] else None
        a["pazarlik"] = yuvarla((a["teklif"] - a["gercek"]) * a["miktar"]) if a["teklif"] else None
        kim = f"{a['ad']} · {a['tedarikci']}".strip(" ·")
        if a["butce"] is None:
            uy.append(uyari("Orta", "Bütçe fiyatı yok", kim, f"{a['gercek_tutar']:,.2f} tutarındaki alımın bütçe fiyatı yazılmamış; harcamaya girdi ama tasarruf hesabına girmedi."))
        elif a["tasarruf"] < 0 and a["oran"] <= -asim_esik:
            uy.append(uyari("Orta", "Bütçe aşımı", kim, f"Bütçe fiyatı {a['butce']:,.2f}, gerçekleşen {a['gercek']:,.2f}; aşım {-a['tasarruf']:,.2f} (%{-a['oran']})."))
        elif a["oran"] >= supheli:
            uy.append(uyari("Bilgi", "Tasarruf oranı çok yüksek", kim, f"Tasarruf oranı %{a['oran']}. Bütçe fiyatının güncel ve aynı özellikteki malzeme için olduğunu doğrulayın."))
        if a["pazarlik"] is not None and a["pazarlik"] < 0:
            uy.append(uyari("Orta", "İlk tekliften pahalı alım", kim, f"İlk teklif {a['teklif']:,.2f}, gerçekleşen {a['gercek']:,.2f}. Kapsam veya özellik değiştiyse not düşün."))
    butceli = [a for a in alimlar if a["butce"]]

    def topla(liste: list[dict]) -> dict:
        b = [a for a in liste if a["butce"]]
        t = {"kalem": len(liste), "harcama": sum((a["gercek_tutar"] for a in liste), Z), "butce": sum((a["butce_tutar"] for a in b), Z), "butceli_harcama": sum((a["gercek_tutar"] for a in b), Z),
             "brut": sum((a["tasarruf"] for a in b if a["tasarruf"] > 0), Z), "asim": sum((-a["tasarruf"] for a in b if a["tasarruf"] < 0), Z), "pazarlik": sum((a["pazarlik"] for a in liste if a["pazarlik"] is not None), Z)}
        t["net"] = t["brut"] - t["asim"]
        t["oran"] = _oran(t["net"], t["butce"])
        return t

    kirilimlar = {}
    for baslik, alan in KIRILIM:
        grup = {}
        for a in alimlar:
            grup.setdefault(a[alan] if alan == "ay" else katla(a[alan]), []).append(a)
        satirlar = [{"ad": g[0][alan], **topla(g)} for g in grup.values()]
        kirilimlar[baslik] = sorted(satirlar, key=lambda s: (s["ad"],) if alan == "ay" else (-s["net"], katla(s["ad"])))
    return {"alimlar": alimlar, "toplam": topla(alimlar), "butcesiz": len(alimlar) - len(butceli), "kirilimlar": kirilimlar, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    t = r["toplam"]
    _baslik(ws, ["Gösterge", "Değer"], (48, 18))
    for etiket, deger in (("Toplam harcama", t["harcama"]), ("Bütçe fiyatı olan alımların bütçe tutarı", t["butce"]), ("Bu alımların gerçekleşen tutarı", t["butceli_harcama"]), ("Brüt tasarruf", t["brut"]), ("Bütçe aşımları", t["asim"]),
                          ("Net tasarruf", t["net"]), ("Net tasarruf oranı %", t["oran"]), ("İlk teklife göre pazarlık kazancı", t["pazarlik"])):
        _satir(ws, [etiket, deger], para_s=(2,))
    _satir(ws, ["Alım sayısı", t["kalem"]])
    _satir(ws, ["Bütçe fiyatı olmayan alım", r["butcesiz"]])
    ws.append([])
    ws.append(["Tasarruf = bütçe tutarı − gerçekleşen tutar. Pazarlık kazancı ayrı bir ölçüdür ve tasarrufa eklenmez."])
    for baslik, satirlar in r["kirilimlar"].items():
        ws = wb.create_sheet(baslik)
        _baslik(ws, [baslik, "Alım", "Harcama", "Bütçe Tutarı", "Brüt Tasarruf", "Aşım", "Net Tasarruf", "Net Tasarruf %", "Pazarlık Kazancı"], (28, 8, 16, 16, 15, 15, 15, 11, 16))
        for s in satirlar:
            _satir(ws, [s["ad"], s["kalem"], s["harcama"], s["butce"], s["brut"] or None, s["asim"] or None, s["net"], s["oran"], s["pazarlik"] or None], para_s=range(3, 10), dolgu={7: KIRMIZI if s["net"] < 0 else None})
    ws = wb.create_sheet("Alımlar")
    _baslik(ws, ["Tarih", "Kategori", "Malzeme / Hizmet", "Tedarikçi", "Miktar", "Bütçe Fiyatı", "İlk Teklif", "Gerçekleşen Fiyat", "Bütçe Tutarı", "Gerçekleşen Tutar", "Tasarruf", "Tasarruf %", "Pazarlık Kazancı", "Tasarruf Türü",
                 "Satın Almacı"], (12, 16, 30, 24, 11, 13, 13, 13, 15, 15, 14, 10, 14, 20, 18))
    for a in r["alimlar"]:
        _satir(ws, [a["tarih"], a["kategori"], a["ad"], a["tedarikci"], a["miktar"], a["butce"], a["teklif"], a["gercek"], a["butce_tutar"], a["gercek_tutar"], a["tasarruf"], a["oran"], a["pazarlik"], a["tur"], a["kisi"]],
               para_s=range(5, 14), tarih_s=(1,), dolgu={11: KIRMIZI if a["tasarruf"] is not None and a["tasarruf"] < 0 else YESIL if a["tasarruf"] else None})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, asim_esik: Decimal = Decimal(5), supheli: Decimal = Decimal(40)) -> dict:
    alimlar, uy = oku(yol)
    r = raporla(alimlar, asim_esik, supheli)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Alımları bütçe fiyatıyla karşılaştırıp satın alma tasarrufunu raporlar.")
    p.add_argument("--alimlar", type=Path, help="Tarih, Kategori, Malzeme, Tedarikçi, Miktar, Bütçe Fiyatı, İlk Teklif Fiyatı, Gerçekleşen Fiyat, Tasarruf Türü, Satın Almacı")
    p.add_argument("--asim", default="5", help="Bütçe aşımı bu yüzde ve üstündeyse uyarılır (varsayılan 5)")
    p.add_argument("--supheli", default="40", help="Tasarruf oranı bu yüzde ve üstündeyse bütçe fiyatı sorgulanır (varsayılan 40)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "satin_alma_tasarruf.xlsx")
    a = p.parse_args(argv)
    yol = a.alimlar or ORNEK / "alimlar.csv"
    asim, supheli = para(a.asim), para(a.supheli)
    if asim is None or supheli is None or asim < 0 or not 0 < supheli <= 100:
        print("[X] --asim sıfır veya pozitif, --supheli 0–100 arasında olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, asim, supheli)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {t['kalem']} alım · harcama {t['harcama']:,.2f} · brüt tasarruf {t['brut']:,.2f} · aşım {t['asim']:,.2f} · net {t['net']:,.2f}" + (f" (%{t['oran']})" if t["oran"] is not None else "")
          + (f" · {r['butcesiz']} alımın bütçe fiyatı yok" if r["butcesiz"] else ""))
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
