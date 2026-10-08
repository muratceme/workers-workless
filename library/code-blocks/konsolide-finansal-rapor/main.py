"""
Konsolide Finansal Rapor — Workers / Workless kod bloğu
Finans › Finans Müdürü

Grup şirketlerinin Tekdüzen mizanlarını birleştirir, grup içi işlemleri eler ve konsolide bilanço / gelir tablosu üretir
(tam konsolidasyon):
  - Grup içi cari bakiyeler (alacak ↔ borç): mutabık tutar elenir, fark bulgu olarak raporlanır ve hesapta kalır.
  - Grup içi satış / hizmet / faiz (gelir ↔ gider) elenir.
  - Grup içi satışlardan stokta kalan gerçekleşmemiş kâr elenir (Dr 621 / Cr 153).
  - Sermaye eliminasyonu: ana ortaklıktaki yatırım (245) ↔ bağlı ortaklık özkaynağı; şerefiye (261), edinim sonrası
    yedeklerin grup payı (570), kontrol gücü olmayan paylar (özkaynak) ve dönem kârındaki payları.
Mali tablo düzeni ve ters bakiye virmanları "Mizandan Mali Tablo Hazırlama" paketiyle aynıdır (ortak çekirdek).
İnternete bağlanmaz.

Kullanım:
    python main.py                                         # örnek: ana ortaklık + 2 bağlı ortaklık
    python main.py --sirketler sirketler.xlsx --eliminasyonlar eliminasyonlar.xlsx --unvan "Örnek Grup"
"""
from __future__ import annotations

import argparse
import re
import sys
from collections import OrderedDict
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import mali_tablo_cekirdek as mt
import mizan_cekirdek as mc

BURASI = Path(__file__).resolve().parent
SIFIR = Decimal(0)
KURUS = Decimal("0.01")
YUZ = Decimal(100)

# Konsolide tabloya özgü satırlar (Tekdüzen'de karşılığı yok): kontrol gücü olmayan paylar özkaynakta ayrı grup,
# dönem kârındaki payları 59 grubunda indirim satırı olarak gösterilir.
KGO, KGO_KAR = "5KG", "59K"
mt.PASIF = mt.PASIF[:-1] + [(mt.PASIF[-1][0], mt.PASIF[-1][1] + [("G", "Kontrol Gücü Olmayan Paylar", ("5K",))])]
OZEL_AD = {"261": "Şerefiye", KGO: "Kontrol Gücü Olmayan Paylar", KGO_KAR: "Kontrol Gücü Olmayan Paylara Ait Dönem Kârı (-)",
           "570": "Geçmiş Yıllar Kârları"}
OZKAYNAK_ONEK = ("50", "52", "54", "57", "58")


def para(x) -> Decimal:
    """mizan çekirdeğinin para()'sı + '750.000' gibi yalnız binlik noktalı yazım."""
    if isinstance(x, str) and re.fullmatch(r"\s*-?\d{1,3}(\.\d{3})+\s*", x):
        x = x.replace(".", "")
    return mc.para(x)


def yuvarla(x: Decimal) -> Decimal:
    return x.quantize(KURUS, rounding=ROUND_HALF_UP)


def _bul(baslik, *adlar):
    b = [mc.katla(x) for x in baslik]
    return next((b.index(mc.katla(a)) for a in adlar if mc.katla(a) in b), None)


def _al(r, i):
    return r[i] if i is not None and i < len(r) else None


def _str(r, i) -> str:
    return str(_al(r, i) or "").strip()


def _tablo(yol: Path, anahtar: str) -> list[list]:
    s = mc.tablo_oku(yol)
    bi = next((i for i, r in enumerate(s[:10]) if any(mc.katla(c) == anahtar for c in r)), 0)
    return s[bi:]


# ----------------------------------------------------------------------------
# Okuma
# ----------------------------------------------------------------------------

def sirketler_oku(yol: Path) -> list[dict]:
    s = _tablo(yol, "mizan")
    b = s[0]
    i = {k: _bul(b, *v) for k, v in {
        "kod": ("kod", "şirket kodu", "şirket"), "unvan": ("ünvan", "unvan", "şirket adı"), "mizan": ("mizan", "mizan dosyası", "dosya"),
        "pay": ("pay %", "pay", "ana ortaklık payı", "ana ortaklık payı %", "sahiplik %", "oy hakkı %"),
        "yatirim": ("yatırım hesabı", "iştirak hesabı", "ana ortaklıktaki hesap"),
        "edinim": ("edinim özkaynağı", "edinim tarihi özkaynağı", "satın alma tarihi özkaynak"), "rol": ("rol", "tür")}.items()}
    if i["kod"] is None or i["mizan"] is None:
        raise SystemExit(f"{yol.name}: Kod ve Mizan sütunları gerekli. Başlıklar: {b}")
    sonuc = []
    for r in s[1:]:
        if not _al(r, i["kod"]):
            continue
        m = Path(_str(r, i["mizan"]))
        sonuc.append({"kod": _str(r, i["kod"]), "unvan": _str(r, i["unvan"]) or _str(r, i["kod"]),
                      "mizan": m if m.is_absolute() else yol.parent / m,
                      "pay": para(_al(r, i["pay"])) if _str(r, i["pay"]) else YUZ,
                      "yatirim": _str(r, i["yatirim"]), "edinim": para(_al(r, i["edinim"])) if _str(r, i["edinim"]) else None,
                      "ana": mc.katla(_al(r, i["rol"])).startswith("ana")})
    if not sonuc:
        raise SystemExit(f"{yol.name}: şirket satırı yok")
    if not any(x["ana"] for x in sonuc):
        sonuc[0]["ana"] = True
    if sum(x["ana"] for x in sonuc) > 1:
        raise SystemExit("Birden fazla ana ortaklık işaretlenmiş")
    return sonuc


def eliminasyon_oku(yol: Path | None) -> list[dict]:
    if not yol:
        return []
    s = _tablo(yol, "tür")
    b = s[0]
    i = {k: _bul(b, *v) for k, v in {
        "tur": ("tür", "tip", "eliminasyon türü"), "sirket": ("şirket", "şirket kodu"), "hesap": ("hesap kodu", "hesap"),
        "karsi": ("karşı şirket", "karşı şirket kodu"), "karsi_hesap": ("karşı hesap kodu", "karşı hesap"),
        "tutar": ("tutar",), "aciklama": ("açıklama",)}.items()}
    if i["tur"] is None or i["sirket"] is None:
        raise SystemExit(f"{yol.name}: Tür ve Şirket sütunları gerekli. Başlıklar: {b}")
    return [{"no": n, "tur": mc.katla(_al(r, i["tur"])), "tur_ham": _str(r, i["tur"]), "sirket": _str(r, i["sirket"]),
             "hesap": _str(r, i["hesap"]).replace(" ", ""), "karsi": _str(r, i["karsi"]), "karsi_hesap": _str(r, i["karsi_hesap"]).replace(" ", ""),
             "tutar": para(_al(r, i["tutar"])) if _str(r, i["tutar"]) else None, "aciklama": _str(r, i["aciklama"])}
            for n, r in enumerate((r for r in s[1:] if _al(r, i["tur"])), 1)]


# ----------------------------------------------------------------------------
# Konsolidasyon
# ----------------------------------------------------------------------------

def bakiye(h: dict, ana: dict, kod: str) -> Decimal | None:
    """Alt hesap (120.02) mizanda varsa onun, 3 haneli ise ana hesabın, yoksa alt kırılımların net bakiyesi."""
    if kod in h:
        return h[kod]["net"]
    if kod in ana:
        return ana[kod]["net"]
    alt = [v["net"] for k, v in h.items() if k.startswith(kod + ".")]
    return sum(alt, SIFIR) if alt else None


def ozkaynak(ana: dict) -> tuple[Decimal, Decimal]:
    """(dönem kârı hariç özkaynak, dönem kârı) — alacak bakiyesi pozitif."""
    oz = -sum((v["net"] for k, v in ana.items() if k.startswith(OZKAYNAK_ONEK)), SIFIR)
    acik_6 = any(v["net"] for k, v in ana.items() if k[0] == "6" and not k.startswith(("690", "692")))
    if acik_6:
        _, kar = mt.gelir_tablosu(ana)
    else:
        kar = -sum((v["net"] for k, v in ana.items() if k.startswith("59")), SIFIR)
    return oz, kar


def konsolide_et(sirketler, veriler, eliminasyonlar, tolerans: Decimal):
    kayitlar, bulgular = [], []
    no = [0]

    def kayit(tur, aciklama, sirket, satirlar):
        """satirlar: [(kod, borç, alacak)]"""
        no[0] += 1
        for kod, b, a in satirlar:
            if b or a:
                kayitlar.append({"no": no[0], "tur": tur, "aciklama": aciklama, "sirket": sirket, "kod": kod[:3] if kod[0].isdigit() else kod,
                                 "alt": kod, "borc": b, "alacak": a})

    def bulgu(seviye, ref, kontrol, aciklama, tutar=None):
        bulgular.append({"seviye": seviye, "ref": ref, "kontrol": kontrol, "aciklama": aciklama, "tutar": tutar})

    def veri(kod):
        if kod not in veriler:
            raise SystemExit(f"Eliminasyon dosyasında tanımsız şirket kodu: {kod}")
        return veriler[kod]

    for e in eliminasyonlar:
        ref = f"#{e['no']} {e['sirket']}/{e['hesap']}" + (f" ↔ {e['karsi']}/{e['karsi_hesap']}" if e["karsi"] else "")
        acik = e["aciklama"] or e["tur_ham"]
        if e["tur"].startswith("cari"):
            a, b = veri(e["sirket"]), veri(e["karsi"])
            x, y = bakiye(a["h"], a["ana"], e["hesap"]), bakiye(b["h"], b["ana"], e["karsi_hesap"])
            if x is None or y is None:
                bulgu("Hata", ref, "Hesap bulunamadı", f"{e['sirket'] if x is None else e['karsi']} mizanında hesap yok; elenmedi")
                continue
            if x == 0 and y == 0:
                continue
            if x * y > 0:
                bulgu("Hata", ref, "Cari yön uyuşmuyor", f"iki taraf da {'borç' if x > 0 else 'alacak'} bakiyeli ({mc.tl(x)} / {mc.tl(y)}); elenmedi")
                continue
            elim = min(abs(x), abs(y))
            isaret = 1 if x > 0 else -1          # x borç bakiyeli (alacak hesabı) ise x alacaklanır, y borçlanır
            kayit("Cari", acik, f"{e['sirket']} ↔ {e['karsi']}",
                  [(e["hesap"], SIFIR, elim) if isaret > 0 else (e["hesap"], elim, SIFIR),
                   (e["karsi_hesap"], elim, SIFIR) if isaret > 0 else (e["karsi_hesap"], SIFIR, elim)])
            fark = abs(x) - abs(y)
            if abs(fark) > tolerans:
                fazla = e["sirket"] if fark > 0 else e["karsi"]
                bulgu("Hata", ref, "Mutabakat farkı", f"{e['sirket']} {mc.tl(abs(x))}, {e['karsi']} {mc.tl(abs(y))}: {mc.tl(abs(fark))} fark "
                      f"{fazla} tarafında elenmeden kaldı (yoldaki fatura/ödeme, kur farkı?)", fark)
        elif e["tur"].startswith(("gelir", "satis", "hizmet", "faiz", "kira")):
            a = veri(e["sirket"])
            veri(e["karsi"])
            gelir = bakiye(a["h"], a["ana"], e["hesap"])
            tutar = e["tutar"] if e["tutar"] is not None else (-gelir if gelir is not None else None)
            if tutar is None:
                bulgu("Hata", ref, "Tutar yok", "gelir hesabı mizanda yok ve tutar verilmemiş; elenmedi")
                continue
            if gelir is not None and e["tutar"] is not None and len(e["hesap"]) > 3 and abs(-gelir - e["tutar"]) > tolerans:
                bulgu("Dikkat", ref, "Gelir hesabı ≠ tutar", f"{e['hesap']} bakiyesi {mc.tl(-gelir)}, mutabakat tutarı {mc.tl(e['tutar'])}",
                      -gelir - e["tutar"])
            kayit("Gelir-Gider", acik, f"{e['sirket']} → {e['karsi']}", [(e["hesap"], tutar, SIFIR), (e["karsi_hesap"] or "621", SIFIR, tutar)])
        elif e["tur"].startswith(("stok", "gerceklesmemis")):
            veri(e["sirket"])
            if not e["tutar"]:
                bulgu("Hata", ref, "Tutar yok", "stoktaki gerçekleşmemiş kâr tutarı verilmemiş")
                continue
            kayit("Stoktaki kâr", acik, e["sirket"], [(e["karsi_hesap"] or "621", e["tutar"], SIFIR), (e["hesap"] or "153", SIFIR, e["tutar"])])
        elif e["tur"].startswith("duzeltme"):
            if not e["tutar"]:
                continue
            kayit("Düzeltme", acik, e["sirket"], [(e["hesap"], e["tutar"], SIFIR), (e["karsi_hesap"], SIFIR, e["tutar"])])
        else:
            bulgu("Hata", ref, "Tanımsız tür", f"'{e['tur_ham']}' (Cari, Gelir-Gider, Stoktaki Kâr, Düzeltme)")

    # Sermaye eliminasyonu
    ana_s = next(s for s in sirketler if s["ana"])
    av = veriler[ana_s["kod"]]
    kgo_toplam = kgo_kar_toplam = serefiye_toplam = SIFIR
    for s in sirketler:
        if s["ana"]:
            continue
        v = veriler[s["kod"]]
        p = s["pay"] / YUZ
        oz, kar = ozkaynak(v["ana"])
        yat_kod = s["yatirim"] or "245"
        yat = bakiye(av["h"], av["ana"], yat_kod)
        if yat is None or yat <= 0:
            bulgu("Hata", s["kod"], "Yatırım hesabı yok", f"ana ortaklık mizanında {yat_kod} bakiyesi yok; sermaye eliminasyonu yapılmadı")
            continue
        if not s["yatirim"] and sum(1 for x in sirketler if not x["ana"]) > 1:
            bulgu("Dikkat", s["kod"], "Yatırım hesabı belirsiz", "birden fazla bağlı ortaklık var; her biri için alt hesap (245.01…) verin")
        edinim = s["edinim"]
        if edinim is None:
            edinim = oz
            bulgu("Dikkat", s["kod"], "Edinim özkaynağı yok", "edinim tarihindeki özkaynak verilmedi: bugünkü özkaynak (dönem kârı hariç) "
                  "kullanıldı; edinimden sonra oluşan yedekler şerefiyeye karışır")
        if not (0 < p <= 1):
            bulgu("Hata", s["kod"], "Pay", f"pay %{s['pay']} geçersiz")
            continue
        if p <= Decimal("0.5"):
            bulgu("Dikkat", s["kod"], "Kontrol", f"pay %{s['pay']:g}: kontrol (TFRS 10) yoksa tam konsolidasyon değil özkaynak yöntemi uygulanır")
        serefiye = yuvarla(yat - p * edinim)
        sonrasi = yuvarla(p * (oz - edinim))
        kgo = yuvarla((1 - p) * oz)
        kgo_kar = yuvarla((1 - p) * kar)
        satir = [(k, -x["net"], SIFIR) if x["net"] < 0 else (k, SIFIR, x["net"])
                 for k, x in sorted(v["ana"].items()) if k.startswith(OZKAYNAK_ONEK) and x["net"]]
        satir.append((yat_kod, SIFIR, yat))
        if serefiye > 0:
            satir.append(("261", serefiye, SIFIR))
        elif serefiye < 0:
            satir.append(("570", SIFIR, -serefiye))
            bulgu("Dikkat", s["kod"], "Negatif şerefiye", f"{mc.tl(-serefiye)}: edinim yılındaysa kâr/zarara (pazarlık satın alım kazancı) "
                  "alınmalı; burada geçmiş yıllar kârlarına yazıldı, gözden geçirin")
        if sonrasi > 0:
            satir.append(("570", SIFIR, sonrasi))
        elif sonrasi < 0:
            satir.append(("570", -sonrasi, SIFIR))
        if kgo:
            satir.append((KGO, SIFIR, kgo) if kgo > 0 else (KGO, -kgo, SIFIR))
        # Yuvarlama farkı → 570
        b_top = sum((b for _, b, _ in satir), SIFIR)
        a_top = sum((a for _, _, a in satir), SIFIR)
        if b_top != a_top:
            satir.append(("570", a_top - b_top, SIFIR) if a_top > b_top else ("570", SIFIR, b_top - a_top))
        kayit("Sermaye", f"{s['unvan']} sermaye eliminasyonu (pay %{s['pay']:g})", f"{ana_s['kod']} ↔ {s['kod']}", satir)
        if kgo_kar:
            kayit("KGO kâr payı", f"{s['unvan']} dönem kârının %{(100 - s['pay']):g}'i kontrol gücü olmayan paylara", s["kod"],
                  [(KGO_KAR, kgo_kar, SIFIR), (KGO, SIFIR, kgo_kar)] if kgo_kar > 0 else [(KGO, -kgo_kar, SIFIR), (KGO_KAR, SIFIR, -kgo_kar)])
        kgo_toplam += kgo + kgo_kar
        kgo_kar_toplam += kgo_kar
        serefiye_toplam += max(serefiye, SIFIR)
    # Birleştir
    toplam: dict[str, dict] = {}
    for s in sirketler:
        for k, x in veriler[s["kod"]]["ana"].items():
            t = toplam.setdefault(k, {"kod": k, "ad": x["ad"] or mc.TDHP_AD.get(k, ""), "net": SIFIR, "borc": None, "alacak": None})
            t["net"] += x["net"]
    konsolide = {k: dict(v) for k, v in toplam.items()}
    for r in kayitlar:
        t = konsolide.setdefault(r["kod"], {"kod": r["kod"], "ad": OZEL_AD.get(r["kod"]) or mc.TDHP_AD.get(r["kod"], ""), "net": SIFIR,
                                            "borc": None, "alacak": None})
        t["net"] += r["borc"] - r["alacak"]
    return konsolide, toplam, kayitlar, bulgular, {"kgo": kgo_toplam, "kgo_kar": kgo_kar_toplam, "serefiye": serefiye_toplam}


def calistir(sirketler_yolu: Path, cikti: Path, eliminasyon_yolu: Path | None = None, unvan: str = "", tolerans: float = 1.0,
             virman: bool = True) -> dict:
    sirketler = sirketler_oku(sirketler_yolu)
    veriler = OrderedDict()
    for s in sirketler:
        h = mc.mizan_oku(s["mizan"])
        ana, virmanlar = mt.virmanla(h) if virman else (mc.ana_hesaplar(h), [])
        t, u = mt.tablolar(ana)
        veriler[s["kod"]] = {"h": h, "ana": ana, "virmanlar": virmanlar, "t": t, "uyarilar": u}
    konsolide, toplam, kayitlar, bulgular, ek = konsolide_et(sirketler, veriler, eliminasyon_oku(eliminasyon_yolu), Decimal(str(tolerans)))
    kt, ku = mt.tablolar(konsolide)
    for s in sirketler:
        for u in veriler[s["kod"]]["uyarilar"]:
            if "kapanış sonrası" not in u:
                bulgular.append({"seviye": "Dikkat", "ref": s["kod"], "kontrol": "Şirket mali tablosu", "aciklama": u, "tutar": None})
    for u in ku:
        if "kapanış sonrası" not in u:
            bulgular.append({"seviye": "Hata" if "denk değil" in u else "Dikkat", "ref": "Konsolide", "kontrol": "Konsolide tablo",
                             "aciklama": u, "tutar": None})
    bulgular.sort(key=lambda b: ({"Hata": 0, "Dikkat": 1, "Bilgi": 2}[b["seviye"]], b["ref"]))
    ana_payi = kt["donem_net"] - ek["kgo_kar"]
    donemler = OrderedDict()
    for s in sirketler:
        v = veriler[s["kod"]]
        donemler[s["kod"]] = {"yol": s["mizan"], "t": v["t"], "uyarilar": v["uyarilar"],
                              "virmanlar": [{**x, "kod": f"{s['kod']}: {x['kod']}"} for x in v["virmanlar"]]}
    donemler["KONSOLİDE"] = {"yol": Path("eliminasyonlar sonrası"), "t": kt, "uyarilar": [], "virmanlar": []}
    mt._rapor(donemler, cikti, unvan)
    _ek_sayfalar(cikti, sirketler, veriler, toplam, konsolide, kayitlar, bulgular, kt, ek, ana_payi, unvan)
    return {"tablo": kt, "konsolide": konsolide, "kayitlar": kayitlar, "bulgular": bulgular, "kgo": ek["kgo"], "kgo_kar": ek["kgo_kar"],
            "serefiye": ek["serefiye"], "ana_payi": ana_payi, "sirket_tablolari": {k: v["t"] for k, v in veriler.items()}}


# ----------------------------------------------------------------------------
# Rapor (çekirdeğin bilanço / gelir tablosu sayfalarına ek)
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Hata": "F8C9C6", "Dikkat": "FFF4CE", "Bilgi": "E8F0FE"}
PARA = '#,##0.00;(#,##0.00);"-"'


def _baslik(ws, satir=1):
    for c in ws[satir]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI


def _ek_sayfalar(cikti, sirketler, veriler, toplam, konsolide, kayitlar, bulgular, kt, ek, ana_payi, unvan):
    wb = load_workbook(cikti)
    o = wb["Bilgi"]
    o["A1"] = f"{unvan or 'Grup'} — Konsolide Mali Tablolar"
    def rol(s):
        return "ana ortaklık" if s["ana"] else f"pay %{s['pay']:g}"
    o.append(["Grup", "; ".join(f"{s['kod']} = {s['unvan']} ({rol(s)})" for s in sirketler)])
    o.append(["Konsolide dönem net kârı", mc.tl(kt["donem_net"])])
    o.append(["  ana ortaklık payı", mc.tl(ana_payi)])
    o.append(["  kontrol gücü olmayan paylar", mc.tl(ek["kgo_kar"])])
    o.append(["Şerefiye", mc.tl(ek["serefiye"])])
    o.append(["Kontrol gücü olmayan paylar (özkaynak)", mc.tl(ek["kgo"])])
    o.append(["Yöntem", "Tam konsolidasyon (TFRS 10 / BOBİ FRS ilkelerine göre basitleştirilmiş): grup içi bakiye ve işlemler "
                        "elenir; şerefiye = yatırım − pay × edinim tarihindeki özkaynak (gerçeğe uygun değer düzeltmeleri, şerefiye "
                        "değer düşüklüğü, ertelenmiş vergi ve yabancı para çevrimi bu araçta yoktur). Tekdüzen hesap düzeninde "
                        "sunulur; yasal konsolide finansal tablolar (TFRS/BOBİ FRS) ve bağımsız denetim yerine geçmez."])
    for r in o.iter_rows():
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    b = wb.create_sheet("Bulgular", 1)
    b.append(["Seviye", "Referans", "Kontrol", "Açıklama", "Tutar"])
    _baslik(b)
    for x in bulgular:
        b.append([x["seviye"], x["ref"], x["kontrol"], x["aciklama"], None if x["tutar"] is None else float(x["tutar"])])
        b.cell(b.max_row, 1).fill = PatternFill("solid", fgColor=RENK[x["seviye"]])
        b.cell(b.max_row, 5).number_format = PARA
    if not bulgular:
        b.append(["", "", "", "Bulgu yok"])
    for j, w in enumerate((9, 34, 24, 90, 15), 1):
        b.column_dimensions[get_column_letter(j)].width = w

    e = wb.create_sheet("Eliminasyon Kayıtları")
    e.append(["No", "Tür", "Açıklama", "Şirket", "Hesap", "Ana Hesap", "Hesap Adı", "Borç", "Alacak"])
    _baslik(e)
    for r in kayitlar:
        e.append([r["no"], r["tur"], r["aciklama"], r["sirket"], r["alt"], r["kod"], konsolide.get(r["kod"], {}).get("ad", ""),
                  float(r["borc"]) or None, float(r["alacak"]) or None])
        for c in (8, 9):
            e.cell(e.max_row, c).number_format = PARA
    e.append(["", "", "TOPLAM", "", "", "", "", float(sum((r["borc"] for r in kayitlar), SIFIR)), float(sum((r["alacak"] for r in kayitlar), SIFIR))])
    for c in e[e.max_row]:
        c.font = Font(bold=True)
        c.number_format = PARA
    for j, w in enumerate((5, 13, 50, 14, 12, 9, 34, 15, 15), 1):
        e.column_dimensions[get_column_letter(j)].width = w
    e.freeze_panes = "C2"

    c = wb.create_sheet("Çalışma Tablosu")
    kodlar = [s["kod"] for s in sirketler]
    c.append(["Hesap", "Hesap Adı"] + kodlar + ["Toplam", "Elim. Borç", "Elim. Alacak", "Konsolide"])
    _baslik(c)
    for k in sorted(konsolide, key=lambda k: (k[:2], k)):
        borc = sum((r["borc"] for r in kayitlar if r["kod"] == k), SIFIR)
        alacak = sum((r["alacak"] for r in kayitlar if r["kod"] == k), SIFIR)
        satir = [veriler[s]["ana"].get(k, {}).get("net") for s in kodlar]
        if not any(satir) and not borc and not alacak and not konsolide[k]["net"]:
            continue
        c.append([k, konsolide[k]["ad"]] + [float(x) if x else None for x in satir]
                 + [float(toplam.get(k, {}).get("net", SIFIR)), float(borc) or None, float(alacak) or None, float(konsolide[k]["net"])])
        for j in range(3, len(kodlar) + 7):
            c.cell(c.max_row, j).number_format = PARA
    c.append(["", "Net bakiye = borç − alacak (borç +, alacak −)"])
    c.column_dimensions["A"].width = 8
    c.column_dimensions["B"].width = 40
    for j in range(3, len(kodlar) + 7):
        c.column_dimensions[get_column_letter(j)].width = 16
    c.freeze_panes = "C2"
    wb.save(cikti)


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ornek = BURASI / "ornek_veri"
    ap = argparse.ArgumentParser(description="Grup şirketlerinin mizanlarını birleştirip grup içi işlemleri eleyerek konsolide tablolar üretir.")
    ap.add_argument("--sirketler", type=Path, default=ornek / "sirketler.csv",
                    help="Kod, Ünvan, Mizan (dosya yolu), Pay %%, Yatırım Hesabı, Edinim Özkaynağı, Rol (Ana/Bağlı)")
    ap.add_argument("--eliminasyonlar", type=Path,
                    help="Tür (Cari / Gelir-Gider / Stoktaki Kâr / Düzeltme), Şirket, Hesap Kodu, Karşı Şirket, Karşı Hesap Kodu, Tutar, Açıklama")
    ap.add_argument("--unvan", default="", help="Grup adı (rapor başlığı)")
    ap.add_argument("--tolerans", type=float, default=1.0, help="Cari mutabakatta kabul edilen fark, TL")
    ap.add_argument("--virman-yok", action="store_true", help="Ters bakiyeli alt hesapları karşı tarafa taşıma")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "konsolide_mali_tablolar.xlsx")
    a = ap.parse_args(argv)
    if a.sirketler == ornek / "sirketler.csv":
        a.eliminasyonlar = a.eliminasyonlar or ornek / "eliminasyonlar.csv"
        a.unvan = a.unvan or "Örnek Grup (kurgusal)"
    s = calistir(a.sirketler, a.cikti, a.eliminasyonlar, a.unvan, a.tolerans, not a.virman_yok)
    t = s["tablo"]
    print(f"[OK] Konsolide aktif {mc.tl(t['aktif_top'])} · pasif {mc.tl(t['pasif_top'])} · net kâr {mc.tl(t['donem_net'])} "
          f"(ana ortaklık {mc.tl(s['ana_payi'])}, KGO {mc.tl(s['kgo_kar'])}) · şerefiye {mc.tl(s['serefiye'])}")
    for x in s["bulgular"]:
        print(f"[{x['seviye']}] {x['ref']}: {x['kontrol']} — {x['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
