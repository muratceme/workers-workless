import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)


def g(n):
    return None if n is None else BUGUN + timedelta(days=n)


def t(no, deger=1000, tur="Çek", gecerlilik=None, ekspertiz=None, sigorta=None, ad="Ad"):
    return {"no": no, "ad": ad, "tur": tur, "deger": D(str(deger)), "gecerlilik": g(gecerlilik), "ekspertiz": ekspertiz, "sigorta": g(sigorta), "nota": ""}


def d(teminatlar, **kw):
    r = main.degerlendir(teminatlar, BUGUN, **kw)
    return {x["no"]: (x["durum"], x["dikkate"], x["sorun"]) for x in r["teminatlar"]}, sorted((u["onem"], u["tur"]) for u in r["uyarilar"])


class TeminatTesti(unittest.TestCase):
    def test_gecerlilik_ve_sigorta(self):
        s, uy = d([t("dün", gecerlilik=-1), t("bugün", gecerlilik=0), t("otuz", gecerlilik=30), t("otuzbir", gecerlilik=31), t("süresiz"), t("sig_dün", sigorta=-1), t("sig_bugün", sigorta=0), t("sig_30", sigorta=30), t("sig_31", sigorta=31)])
        self.assertEqual(s, {"dün": ("Süresi dolmuş", D(0), "Geçerlilik süresi dolmuş"), "bugün": ("Aksiyon gerekli", D(1000), "Geçerlilik 0 gün sonra doluyor"), "otuz": ("Aksiyon gerekli", D(1000), "Geçerlilik 30 gün sonra doluyor"),
                             "otuzbir": ("Geçerli", D(1000), ""), "süresiz": ("Geçerli", D(1000), ""), "sig_dün": ("Aksiyon gerekli", D(1000), "Sigorta süresi dolmuş"), "sig_bugün": ("Aksiyon gerekli", D(1000), "Sigorta bitişi yaklaşıyor"),
                             "sig_30": ("Aksiyon gerekli", D(1000), "Sigorta bitişi yaklaşıyor"), "sig_31": ("Geçerli", D(1000), "")})
        self.assertEqual(uy, [("Orta", "Sigorta bitişi yaklaşıyor")] * 2 + [("Orta", "Teminatın süresi yaklaşıyor")] * 2 + [("Yüksek", "Sigorta süresi dolmuş"), ("Yüksek", "Teminatın süresi dolmuş")])
        self.assertEqual((d([t("a", gecerlilik=45)], gun=45)[0]["a"][0], d([t("a", gecerlilik=46)], gun=45)[0]["a"][0]), ("Aksiyon gerekli", "Geçerli"))

    def test_ekspertiz_ve_oran(self):
        l = [t("dün", tur="İpotek", ekspertiz=date(2025, 10, 11)), t("bugün", tur="İpotek", ekspertiz=date(2025, 10, 12)), t("otuz", tur="İpotek", ekspertiz=date(2025, 11, 11)), t("otuzbir", tur="İpotek", ekspertiz=date(2025, 11, 12)),
             t("yok", tur="1. Derece İPOTEK"), t("çek"), t("rehin", tur="Rehin", ekspertiz=date(2020, 1, 1))]
        s, uy = d(l)
        self.assertEqual({k: v[2] for k, v in s.items()}, {"dün": "Ekspertiz yenilenmeli", "bugün": "Ekspertiz yenileme yaklaşıyor", "otuz": "Ekspertiz yenileme yaklaşıyor", "otuzbir": "", "yok": "Ekspertiz tarihi yok", "çek": "",
                                                           "rehin": "Ekspertiz yenilenmeli"})
        self.assertEqual(uy, [("Bilgi", "Ekspertiz yenileme yaklaşıyor")] * 2 + [("Orta", "Ekspertiz tarihi yok")] + [("Orta", "Ekspertiz yenilenmeli")] * 2)
        r = main.degerlendir([t("a", tur="İpotek", ekspertiz=date(2024, 10, 13))], BUGUN, ekspertiz_ay=24)["teminatlar"][0]
        self.assertEqual((r["ekspertiz_son"], r["sorun"], main.ay_ekle(date(2026, 8, 31), 6)), (date(2026, 10, 13), "Ekspertiz yenileme yaklaşıyor", date(2027, 2, 28)))
        r = main.degerlendir([t("1", 1000, "İpotek", ekspertiz=BUGUN), t("2", "333.33", "ÇEK"), t("3", 500, "Senet"), t("4", 100, "senet"), t("5", 900, "ipotek", gecerlilik=-5, ekspertiz=BUGUN)], BUGUN, oranlar={"ipotek": D(75), "cek": D("72.5")})
        self.assertEqual([(x["oran"], x["dikkate"]) for x in r["teminatlar"]], [(D(75), D("750.00")), (D("72.5"), D("241.66")), (D(100), D("500.00")), (D(100), D("100.00")), (D(75), D(0))])
        self.assertEqual(([(u["tur"], u["kim"]) for u in r["uyarilar"] if u["onem"] == "Orta"], [(x["tur"], x["adet"], x["deger"], x["dikkate"], x["oran"]) for x in r["turler"]]),
                         ([("Dikkate alma oranı tanımsız", "Senet")], [("İpotek", 2, D(1900), D("750.00"), D(75)), ("Senet", 2, D(600), D("600.00"), D(100)), ("ÇEK", 1, D("333.33"), D("241.66"), D("72.5"))]))

    def test_musteri(self):
        tem = [t("M1", 600), t("m1", 300, gecerlilik=-1), t("M2", 1000), t("M3", 500, ad="")]
        risk = [{"no": "M1", "ad": "Bir", "risk": D(800)}, {"no": "M2", "ad": "İki", "risk": D(1000)}, {"no": "M4", "ad": "Dört", "risk": D(250)}, {"no": "M2", "ad": "İki", "risk": D(0)}, {"no": "M5", "ad": "Beş", "risk": D(0)}]
        r = main.degerlendir(tem, BUGUN, risk)
        self.assertEqual([(o["no"], o["ad"], o["teminat"], o["deger"], o["dikkate"], o["risk"], o["karsilama"], o["acik"], o["sorunlu"]) for o in r["musteriler"]],
                         [("M4", "Dört", 0, D(0), D(0), D(250), D("0.00"), D(250), 0), ("M1", "Bir", 2, D(900), D(600), D(800), D("75.00"), D(200), 1), ("M2", "İki", 1, D(1000), D(1000), D(1000), D("100.00"), D(0), 0),
                          ("M3", "", 1, D(500), D(500), None, None, None, 0), ("M5", "Beş", 0, D(0), D(0), D(0), None, D(0), 0)])
        self.assertEqual((r["toplam"], r["risk_var"], sorted(u["kim"] for u in r["uyarilar"] if u["tur"] == "Teminat açığı")), ({"deger": D(2400), "dikkate": D(2100), "risk": D(2050), "acik": D(450)}, True, ["M1 · Bir", "M4 · Dört"]))
        self.assertEqual(sorted(u["kim"] for u in main.degerlendir(tem, BUGUN, risk, karsilama=D(120))["uyarilar"] if u["tur"] == "Teminat açığı"), ["M1 · Bir", "M2 · İki", "M4 · Dört"])
        r = main.degerlendir(tem, BUGUN)
        self.assertEqual((r["risk_var"], r["toplam"]["risk"], [o["karsilama"] for o in r["musteriler"]], r["say"]), (False, D(0), [None] * 3, {"Süresi dolmuş": 1, "Aksiyon gerekli": 0, "Geçerli": 3}))
        self.assertEqual(len(main.degerlendir([], BUGUN, risk)["musteriler"]), 4)
        with self.assertRaises(ValueError):
            main.degerlendir([], BUGUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "teminatlar.csv", cikti, BUGUN, ORNEK / "riskler.csv", ORNEK / "oranlar.csv")
            self.assertEqual([(x["no"], x["tur"], x["durum"], x["dikkate"]) for x in r["teminatlar"]],
                             [("2001", "İpotek", "Geçerli", D(3750000)), ("2001", "Müşteri Çeki", "Aksiyon gerekli", D(320000)), ("2002", "Teminat Mektubu", "Süresi dolmuş", D(0)), ("2002", "Kefalet", "Geçerli", D(500000)),
                              ("2003", "İpotek", "Aksiyon gerekli", D(2250000)), ("2003", "Araç Rehni", "Aksiyon gerekli", D(360000)), ("2004", "İpotek", "Aksiyon gerekli", D(1500000)), ("2005", "Nakit Blokaj", "Geçerli", D(800000))])
            self.assertEqual([(o["no"], o["dikkate"], o["risk"], o["karsilama"], o["acik"]) for o in r["musteriler"]],
                             [("2002", D(500000), D(1200000), D("41.67"), D(700000)), ("2006", D(0), D(500000), D("0.00"), D(500000)), ("2005", D(800000), D(900000), D("88.89"), D(100000)), ("2001", D(4070000), D(3500000), D("116.29"), D(0)),
                              ("2003", D(2610000), D(2400000), D("108.75"), D(0)), ("2004", D(1500000), D(1000000), D("150.00"), D(0))])
            self.assertEqual((r["toplam"], r["say"]), ({"deger": D(14300000), "dikkate": D(9480000), "risk": D(9500000), "acik": D(1300000)}, {"Süresi dolmuş": 1, "Aksiyon gerekli": 4, "Geçerli": 3}))
            self.assertEqual([(x["tur"], x["adet"], x["deger"], x["dikkate"]) for x in r["turler"]],
                             [("İpotek", 3, D(10000000), D(7500000)), ("Teminat Mektubu", 1, D(1500000), D(0)), ("Kefalet", 1, D(1000000), D(500000)), ("Nakit Blokaj", 1, D(800000), D(800000)), ("Araç Rehni", 1, D(600000), D(360000)),
                              ("Müşteri Çeki", 1, D(400000), D(320000))])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Yüksek", "Sigorta süresi dolmuş", "2003 · Gama Gıda · İpotek"), ("Yüksek", "Teminatın süresi dolmuş", "2002 · Beta Tekstil · Teminat Mektubu"), ("Orta", "Ekspertiz tarihi yok", "2004 · Delta Lojistik · İpotek"),
                              ("Orta", "Ekspertiz yenilenmeli", "2003 · Gama Gıda · İpotek"), ("Orta", "Sigorta bitişi yaklaşıyor", "2003 · Gama Gıda · Araç Rehni"), ("Orta", "Teminat açığı", "2002 · Beta Tekstil"),
                              ("Orta", "Teminat açığı", "2005 · Epsilon Kimya"), ("Orta", "Teminat açığı", "2006 · Zeta Mobilya"), ("Orta", "Teminatın süresi yaklaşıyor", "2001 · Alfa İnşaat · Müşteri Çeki")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Teminatlar", "Müşteriler", "Teminat Türleri", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            ty, ry, oy = Path(tmp) / "t.csv", Path(tmp) / "r.csv", Path(tmp) / "o.csv"
            ty.write_text("No;Unvan;Teminat;Tutar;Vade;Değerleme Tarihi;Poliçe Bitiş\n7;A Ltd.;İpotek;1.250.000,50;x;01.02.2026;01.03.2027\n8;;Çek;x;;;\n9;;;5;;;\n;B;Çek;5;;;\n", encoding="utf-8")
            ry.write_text("No;Risk\n7;900.000\n;5\n", encoding="utf-8")
            oy.write_text("Teminat;Marj\nİpotek;72,5\n;5\n", encoding="utf-8")
            l, uy = main.oku_teminatlar(ty)
            self.assertEqual(([(x["no"], x["ad"], x["tur"], x["deger"], x["gecerlilik"], x["ekspertiz"], x["sigorta"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("7", "A Ltd.", "İpotek", D("1250000.50"), None, date(2026, 2, 1), date(2027, 3, 1))], [("Tarih okunamadı", "7 · A Ltd. · İpotek"), ("Teminat satırı okunamadı", "8 · Çek"), ("Teminat satırı okunamadı", "9")]))
            self.assertEqual((main.oku_riskler(ry), main.oku_oranlar(oy)), ([{"no": "7", "ad": "", "risk": D(900000)}], {"ipotek": D("72.5")}))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--teminatlar", str(ty), "--cikti", c]
            self.assertEqual(main.main(arg + ["--riskler", str(ry), "--oranlar", str(oy), "--tarih", "12.10.2026", "--gun", "45", "--ekspertiz-ay", "24", "--karsilama", "120"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--gun", "-1"], ["--ekspertiz-ay", "0"], ["--karsilama", "0"])], [2] * 4)
            self.assertEqual(main.main(arg + ["--riskler", str(Path(tmp) / "yok.csv")]), 1)
            ry.write_text("Müşteri No;Kredi Riski\n7;x\n", encoding="utf-8")
            oy.write_text("Teminat Türü;Oran\nİpotek;101\n", encoding="utf-8")
            self.assertEqual([main.main(arg + ["--riskler", str(ry)]), main.main(arg + ["--oranlar", str(oy)])], [1, 1])
            for icerik in ("Müşteri No;Teminat Türü\nA;B\n", "Müşteri No;Teminat Türü;Teminat Değeri\n;B;1\n"):
                ty.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
