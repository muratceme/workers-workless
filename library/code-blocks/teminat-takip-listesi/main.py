"""
Teminat Takip Listesi — Workers / Workless kod bloğu
Şube Bankacılığı › Ticari Portföy Yöneticisi

Kredi teminatlarının (ipotek, çek, kefalet, teminat mektubu) değer, vade ve ekspertiz yenileme tarihlerini izler.
  Dikkate alınan değer = teminat değeri × dikkate alma oranı   (oranlar dosyasından; dosya yoksa %100)
                         geçerlilik tarihi geçmiş teminat sıfır alınır
  Karşılama %          = müşterinin dikkate alınan teminat toplamı ÷ kredi riski × 100
  Teminat açığı        = kredi riski − dikkate alınan teminat toplamı   (eksiye düşmez)
  Ekspertiz yenileme   = ekspertiz tarihi + --ekspertiz-ay (12) ay; geçmişse "yenilenmeli"
  Geçerlilik / sigorta = bitiş tarihi geçmişse "dolmuş", --gun (30) gün veya az kaldıysa "yaklaşıyor"
Dikkate alma oranları ve ekspertiz yenileme süresi bankanızın kredi politikasıdır; kod bu değerleri bilmez, sizden alır.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
import calendar

T_SUTUN = {"no": ("musteri no", "musteri numarasi", "no"), "ad": ("musteri", "musteri adi", "unvan"), "tur": ("teminat turu", "teminat", "tur"), "deger": ("teminat degeri", "deger", "tutar"),
           "gecerlilik": ("gecerlilik tarihi", "vade", "vade tarihi", "bitis tarihi"), "ekspertiz": ("ekspertiz tarihi", "degerleme tarihi", "ekspertiz"), "sigorta": ("sigorta bitis", "sigorta bitis tarihi", "police bitis"),
           "nota": ("aciklama", "not")}
R_SUTUN = {"no": ("musteri no", "musteri numarasi", "no"), "ad": ("musteri", "musteri adi", "unvan"), "risk": ("kredi riski", "risk", "risk bakiyesi", "toplam risk")}
O_SUTUN = {"tur": ("teminat turu", "teminat", "tur"), "oran": ("dikkate alma orani", "dikkate alma orani %", "oran", "marj")}
DURUM_RENK = {"Süresi dolmuş": KIRMIZI, "Aksiyon gerekli": SARI, "Geçerli": YESIL}
Z, YUZ = Decimal(0), Decimal(100)


def ay_ekle(t: date, n: int) -> date:
    y, a = divmod(t.year * 12 + t.month - 1 + n, 12)
    return date(y, a + 1, min(t.day, calendar.monthrange(y, a + 1)[1]))


def olcu(x) -> Decimal | None:
    """Oran okur: virgül ya da tek nokta ondalık ayırıcıdır (72,5 = 72.5)."""
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x or ""))
    try:
        return Decimal(s.replace(".", "").replace(",", ".") if "," in s else s) if s else None
    except InvalidOperation:
        return None


def oku_teminatlar(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, T_SUTUN, ("no", "tur", "deger")):
        no = metin(r["no"])
        if not no:
            continue
        kim, deger = f"{no} · {metin(r['ad'])} · {metin(r['tur'])}".replace(" ·  ·", " ·").strip(" ·"), para(r["deger"])
        if not metin(r["tur"]) or deger is None or deger < 0:
            uy.append(uyari("Yüksek", "Teminat satırı okunamadı", kim, f"Satır {r['_satir']}: teminat türü '{metin(r['tur'])}', değer '{metin(r['deger'])}'. Satır alınmadı."))
            continue
        for alan, ad in (("gecerlilik", "Geçerlilik"), ("ekspertiz", "Ekspertiz"), ("sigorta", "Sigorta bitiş")):
            if metin(r[alan]) and tarih(r[alan]) is None:
                uy.append(uyari("Orta", "Tarih okunamadı", kim, f"{ad} tarihi '{metin(r[alan])}' okunamadı; boş sayıldı."))
        l.append({"no": no, "ad": metin(r["ad"]), "tur": metin(r["tur"]), "deger": deger, "gecerlilik": tarih(r["gecerlilik"]), "ekspertiz": tarih(r["ekspertiz"]), "sigorta": tarih(r["sigorta"]), "nota": metin(r["nota"])})
    return l, uy


def oku_riskler(yol: Path) -> list[dict]:
    l = []
    for r in kayitlar(yol, R_SUTUN, ("no", "risk")):
        if not metin(r["no"]):
            continue
        risk = para(r["risk"])
        if risk is None or risk < 0:
            raise ValueError(f"{yol.name} satır {r['_satir']}: kredi riski '{metin(r['risk'])}' sıfır veya pozitif bir sayı olmalı.")
        l.append({"no": metin(r["no"]), "ad": metin(r["ad"]), "risk": risk})
    return l


def oku_oranlar(yol: Path) -> dict[str, Decimal]:
    oranlar = {}
    for r in kayitlar(yol, O_SUTUN, ("tur", "oran")):
        if not metin(r["tur"]):
            continue
        oran = olcu(r["oran"])
        if oran is None or not 0 <= oran <= 100:
            raise ValueError(f"{yol.name} satır {r['_satir']}: dikkate alma oranı '{metin(r['oran'])}' 0–100 arasında olmalı.")
        oranlar[katla(r["tur"])] = oran
    return oranlar


def degerlendir(teminatlar: list[dict], bugun: date, riskler: list[dict] | None = None, oranlar: dict[str, Decimal] | None = None, gun: int = 30, ekspertiz_ay: int = 12, karsilama: Decimal = YUZ) -> dict:
    if not teminatlar and not riskler:
        raise ValueError("Teminat ve risk listesi boş.")
    uy, musteriler, tanimsiz = [], {}, set()
    for r in riskler or []:
        m = musteriler.setdefault(katla(r["no"]), {"no": r["no"], "ad": r["ad"], "risk": Z, "teminatlar": []})
        m["risk"] += r["risk"]
    for t in teminatlar:
        kim, sorunlar = f"{t['no']} · {t['ad']} · {t['tur']}".replace(" ·  ·", " ·"), []
        oran = (oranlar or {}).get(katla(t["tur"]))
        if oranlar is not None and oran is None and katla(t["tur"]) not in tanimsiz:
            tanimsiz.add(katla(t["tur"]))
            uy.append(uyari("Orta", "Dikkate alma oranı tanımsız", t["tur"], "Bu teminat türü oran dosyasında yok; %100 alındı."))
        t["oran"] = oran if oran is not None else YUZ
        t["kalan"] = (t["gecerlilik"] - bugun).days if t["gecerlilik"] else None
        dolmus = t["kalan"] is not None and t["kalan"] < 0
        t["dikkate"] = Z if dolmus else yuvarla(t["deger"] * t["oran"] / YUZ)
        if dolmus:
            sorunlar.append("Geçerlilik süresi dolmuş")
            uy.append(uyari("Yüksek", "Teminatın süresi dolmuş", kim, f"{t['deger']:,.2f} değerindeki teminatın geçerliliği {t['gecerlilik']:%d.%m.%Y} tarihinde doldu ({-t['kalan']} gün önce). Dikkate alınan değer sıfırlandı."))
        elif t["kalan"] is not None and t["kalan"] <= gun:
            sorunlar.append(f"Geçerlilik {t['kalan']} gün sonra doluyor")
            uy.append(uyari("Orta", "Teminatın süresi yaklaşıyor", kim, f"{t['deger']:,.2f} değerindeki teminatın geçerliliği {t['gecerlilik']:%d.%m.%Y} tarihinde doluyor ({t['kalan']} gün kaldı)."))
        t["ekspertiz_son"] = ay_ekle(t["ekspertiz"], ekspertiz_ay) if t["ekspertiz"] else None
        if t["ekspertiz_son"] and t["ekspertiz_son"] < bugun:
            sorunlar.append("Ekspertiz yenilenmeli")
            uy.append(uyari("Orta", "Ekspertiz yenilenmeli", kim, f"Son ekspertiz {t['ekspertiz']:%d.%m.%Y}; {ekspertiz_ay} aylık süre {t['ekspertiz_son']:%d.%m.%Y} tarihinde doldu."))
        elif t["ekspertiz_son"] and (t["ekspertiz_son"] - bugun).days <= gun:
            sorunlar.append("Ekspertiz yenileme yaklaşıyor")
            uy.append(uyari("Bilgi", "Ekspertiz yenileme yaklaşıyor", kim, f"Son ekspertiz {t['ekspertiz']:%d.%m.%Y}; yenileme tarihi {t['ekspertiz_son']:%d.%m.%Y}."))
        elif t["ekspertiz"] is None and "ipotek" in katla(t["tur"]):
            sorunlar.append("Ekspertiz tarihi yok")
            uy.append(uyari("Orta", "Ekspertiz tarihi yok", kim, "İpotek için ekspertiz tarihi yazılmamış; değerin güncelliği izlenemiyor."))
        if t["sigorta"] and t["sigorta"] < bugun:
            sorunlar.append("Sigorta süresi dolmuş")
            uy.append(uyari("Yüksek", "Sigorta süresi dolmuş", kim, f"Teminatın sigortası {t['sigorta']:%d.%m.%Y} tarihinde bitti ({(bugun - t['sigorta']).days} gün önce)."))
        elif t["sigorta"] and (t["sigorta"] - bugun).days <= gun:
            sorunlar.append("Sigorta bitişi yaklaşıyor")
            uy.append(uyari("Orta", "Sigorta bitişi yaklaşıyor", kim, f"Teminatın sigortası {t['sigorta']:%d.%m.%Y} tarihinde bitiyor ({(t['sigorta'] - bugun).days} gün kaldı)."))
        t["durum"], t["sorun"] = "Süresi dolmuş" if dolmus else "Aksiyon gerekli" if sorunlar else "Geçerli", "; ".join(sorunlar)
        m = musteriler.setdefault(katla(t["no"]), {"no": t["no"], "ad": t["ad"], "risk": None, "teminatlar": []})
        m["ad"] = m["ad"] or t["ad"]
        m["teminatlar"].append(t)
    ozet = []
    for m in musteriler.values():
        o = {"no": m["no"], "ad": m["ad"], "teminat": len(m["teminatlar"]), "deger": sum((t["deger"] for t in m["teminatlar"]), Z), "dikkate": sum((t["dikkate"] for t in m["teminatlar"]), Z), "risk": m["risk"],
             "sorunlu": sum(1 for t in m["teminatlar"] if t["durum"] != "Geçerli")}
        o["karsilama"] = yuvarla(o["dikkate"] * YUZ / o["risk"]) if o["risk"] else None
        o["acik"] = max(o["risk"] - o["dikkate"], Z) if o["risk"] is not None else None
        if o["karsilama"] is not None and o["karsilama"] < karsilama:
            uy.append(uyari("Orta", "Teminat açığı", f"{o['no']} · {o['ad']}".strip(" ·"), f"Kredi riski {o['risk']:,.2f}, dikkate alınan teminat {o['dikkate']:,.2f} (%{o['karsilama']}; hedef %{karsilama:f}). Açık {o['acik']:,.2f}."))
        ozet.append(o)
    ozet.sort(key=lambda o: (-(o["acik"] or Z), katla(o["no"])))
    turler = {}
    for t in teminatlar:
        g = turler.setdefault(katla(t["tur"]), {"tur": t["tur"], "adet": 0, "deger": Z, "dikkate": Z, "oran": t["oran"]})
        g["adet"] += 1
        g["deger"] += t["deger"]
        g["dikkate"] += t["dikkate"]
    riskli = [o for o in ozet if o["risk"] is not None]
    return {"teminatlar": teminatlar, "musteriler": ozet, "turler": sorted(turler.values(), key=lambda g: (-g["deger"], katla(g["tur"]))), "risk_var": riskler is not None,
            "toplam": {"deger": sum((t["deger"] for t in teminatlar), Z), "dikkate": sum((t["dikkate"] for t in teminatlar), Z), "risk": sum((o["risk"] for o in riskli), Z), "acik": sum((o["acik"] for o in riskli), Z)},
            "say": {d: sum(1 for t in teminatlar if t["durum"] == d) for d in DURUM_RENK}, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, bugun: date) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Teminatlar"
    _baslik(ws, ["Durum", "Müşteri No", "Müşteri", "Teminat Türü", "Teminat Değeri", "Oran %", "Dikkate Alınan", "Geçerlilik", "Kalan Gün", "Ekspertiz", "Ekspertiz Yenileme", "Sigorta Bitiş", "Yapılacak", "Açıklama"],
            (16, 12, 26, 20, 16, 8, 16, 12, 9, 12, 13, 12, 44, 30))
    for t in sorted(r["teminatlar"], key=lambda t: (list(DURUM_RENK).index(t["durum"]), katla(t["no"]))):
        _satir(ws, [t["durum"], t["no"], t["ad"], t["tur"], t["deger"], t["oran"], t["dikkate"], t["gecerlilik"], t["kalan"], t["ekspertiz"], t["ekspertiz_son"], t["sigorta"], t["sorun"], t["nota"]], para_s=(5, 6, 7), tarih_s=(8, 10, 11, 12),
               dolgu={1: DURUM_RENK[t["durum"]]}, sar=(13, 14))
    ws = wb.create_sheet("Müşteriler")
    _baslik(ws, ["Müşteri No", "Müşteri", "Teminat Sayısı", "Teminat Değeri", "Dikkate Alınan", "Kredi Riski", "Karşılama %", "Teminat Açığı", "Sorunlu Teminat"], (12, 28, 9, 16, 16, 16, 11, 16, 10))
    for o in r["musteriler"]:
        _satir(ws, [o["no"], o["ad"], o["teminat"], o["deger"], o["dikkate"], o["risk"], o["karsilama"], o["acik"] or None, o["sorunlu"] or None], para_s=(4, 5, 6, 7, 8), dolgu={8: KIRMIZI if o["acik"] else None})
    t = r["toplam"]
    n = _satir(ws, ["TOPLAM", None, len(r["teminatlar"]), t["deger"], t["dikkate"], t["risk"] if r["risk_var"] else None, None, t["acik"] if r["risk_var"] else None, None], para_s=(4, 5, 6, 8))
    for h in ws[n]:
        h.font = Font(bold=True)
    ws.append([])
    ws.append([f"Durum tarihi: {bugun:%d.%m.%Y}. Dikkate alınan değer = teminat değeri × oran; süresi dolmuş teminat sıfır alınır."])
    ws = wb.create_sheet("Teminat Türleri")
    _baslik(ws, ["Teminat Türü", "Adet", "Teminat Değeri", "Oran %", "Dikkate Alınan"], (26, 8, 16, 9, 16))
    for g in r["turler"]:
        _satir(ws, [g["tur"], g["adet"], g["deger"], g["oran"], g["dikkate"]], para_s=(3, 4, 5))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(t_yol: Path, cikti: Path, bugun: date, r_yol: Path | None = None, o_yol: Path | None = None, gun: int = 30, ekspertiz_ay: int = 12, karsilama: Decimal = YUZ) -> dict:
    teminatlar, uy = oku_teminatlar(t_yol)
    r = degerlendir(teminatlar, bugun, oku_riskler(r_yol) if r_yol else None, oku_oranlar(o_yol) if o_yol else None, gun, ekspertiz_ay, karsilama)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, bugun)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Kredi teminatlarının değer, vade, ekspertiz ve sigorta tarihlerini izler.")
    p.add_argument("--teminatlar", type=Path, help="Müşteri No, Müşteri, Teminat Türü, Teminat Değeri, Geçerlilik Tarihi, Ekspertiz Tarihi, Sigorta Bitiş, Açıklama")
    p.add_argument("--riskler", type=Path, help="Müşteri No, Müşteri, Kredi Riski (isteğe bağlı; karşılama oranı için)")
    p.add_argument("--oranlar", type=Path, help="Teminat Türü, Dikkate Alma Oranı (isteğe bağlı; verilmezse hepsi yüzde 100)")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--gun", type=int, default=30, help="Geçerlilik, sigorta ve ekspertiz için kaç gün önceden uyarılsın (varsayılan 30)")
    p.add_argument("--ekspertiz-ay", type=int, default=12, help="Ekspertizin kaç ayda bir yenileneceği (varsayılan 12; banka politikanıza göre değiştirin)")
    p.add_argument("--karsilama", default="100", help="Karşılama oranı bu yüzdenin altındaysa teminat açığı uyarılır (varsayılan 100)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "teminat_takip_listesi.xlsx")
    a = p.parse_args(argv)
    ornek = a.teminatlar is None
    t_yol = ORNEK / "teminatlar.csv" if ornek else a.teminatlar
    r_yol = ORNEK / "riskler.csv" if ornek and a.riskler is None else a.riskler
    o_yol = ORNEK / "oranlar.csv" if ornek and a.oranlar is None else a.oranlar
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    karsilama = para(a.karsilama)
    if bugun is None or not 0 <= a.gun <= 365 or not 1 <= a.ekspertiz_ay <= 120 or karsilama is None or karsilama <= 0:
        print("[X] --tarih GG.AA.YYYY; --gun 0–365; --ekspertiz-ay 1–120; --karsilama pozitif olmalı")
        return 2
    for y in (t_yol, r_yol, o_yol):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(t_yol, a.cikti, bugun, r_yol, o_yol, a.gun, a.ekspertiz_ay, karsilama)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t, s = r["toplam"], r["say"]
    print(f"[OK] {len(r['teminatlar'])} teminat, {len(r['musteriler'])} müşteri · değer {t['deger']:,.2f} · dikkate alınan {t['dikkate']:,.2f}" + (f" · risk {t['risk']:,.2f} · açık {t['acik']:,.2f}" if r["risk_var"] else "")
          + f" · süresi dolmuş {s['Süresi dolmuş']}, aksiyon gerekli {s['Aksiyon gerekli']}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
