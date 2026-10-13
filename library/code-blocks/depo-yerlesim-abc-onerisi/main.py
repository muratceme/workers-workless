"""
Depo Yerleşim (ABC) Önerisi — Workers / Workless kod bloğu
Lojistik ve Depo › Depo Sorumlusu

Ürünleri çıkış sıklığına göre ABC sınıflarına ayırır; sık çıkanların sevkiyata yakın bölgede durmasını önerir.
  Çıkış sıklığı    = ürünün çıkış (toplama) satırı sayısı   (miktar değil; rafa kaç kez gidildiği)
  Pay %            = ürünün sıklığı ÷ toplam sıklık × 100;   kümülatif % = sıklığa göre azalan sırada birikimli pay
  Sınıf            = A: kendinden önceki kümülatif pay %80'in altında · B: %95'in altında · C: kalan · D: dönemde hiç çıkışı yok
  Önerilen bölge   = A sınıfı → A bölgesi (sevkiyata en yakın) · B → B · C ve D → C
  Öneri            = Öne al (mevcut bölge önerilenden uzak) · Geriye al (mevcut bölge önerilenden yakın) · Yerinde
Bölgeleri (A/B/C) deponuza göre siz tanımlarsınız. Kod ürün boyutunu, ağırlığını, raf kapasitesini ve birlikte çıkan ürünleri bilmez.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
C_SUTUN = {"tarih": ("tarih", "cikis tarihi", "donem"), "kod": ("stok kodu", "urun kodu", "malzeme kodu", "kod"), "sayi": ("cikis sayisi", "siklik", "satir sayisi", "toplama sayisi"), "miktar": ("miktar", "cikis miktari")}
U_SUTUN = {"kod": ("stok kodu", "urun kodu", "malzeme kodu", "kod"), "ad": ("stok adi", "urun", "urun adi", "malzeme", "ad"), "konum": ("mevcut konum", "konum", "raf", "adres"), "bolge": ("bolge", "mevcut bolge", "zon")}
RANK = {"A": 0, "B": 1, "C": 2}
ONERI_RENK = {"Öne al": SARI, "Geriye al": MAVI, "Yerinde": YESIL, "Bölge bilinmiyor": None}
Z, YUZ = Decimal(0), Decimal(100)


def olcu(x) -> Decimal | None:
    """Miktar okur: virgül ya da tek nokta ondalık ayırıcıdır; '1.250,5' binlik ayırıcılıdır."""
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x or ""))
    try:
        return Decimal(s.replace(".", "").replace(",", ".") if "," in s else s) if s else None
    except InvalidOperation:
        return None


def oku_cikislar(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, C_SUTUN, ("kod",)):
        kod = metin(r["kod"])
        if not kod:
            continue
        sayi, miktar = (para(r["sayi"]) if metin(r["sayi"]) else Decimal(1)), olcu(r["miktar"])
        if sayi is None or sayi < 0 or sayi != sayi.to_integral_value() or (miktar is not None and miktar < 0):
            uy.append(uyari("Yüksek", "Çıkış satırı okunamadı", kod, f"Satır {r['_satir']}: çıkış sayısı '{metin(r['sayi'])}' (sıfır veya pozitif tam sayı olmalı), miktar '{metin(r['miktar'])}'. Satır alınmadı."))
            continue
        l.append({"tarih": tarih(r["tarih"]), "kod": kod, "sayi": int(sayi), "miktar": miktar or Z})
    return l, uy


def oku_urunler(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, U_SUTUN, ("kod",)):
        kod = metin(r["kod"])
        if not kod:
            continue
        bolge = " ".join(k for k in katla(r["bolge"]).split() if k not in ("bolge", "bolgesi", "zon")).upper()
        if bolge and bolge not in RANK:
            uy.append(uyari("Orta", "Bölge tanımsız", kod, f"Bölge '{metin(r['bolge'])}' okunamadı; A, B ya da C yazın. Boş sayıldı."))
            bolge = ""
        l.append({"kod": kod, "ad": metin(r["ad"]), "konum": metin(r["konum"]), "bolge": bolge})
    return l, uy


def onerile(cikislar: list[dict], urunler: list[dict] | None = None, a: Decimal = Decimal(80), b: Decimal = Decimal(95)) -> dict:
    uy, u = [], {}
    for x in urunler or []:
        if katla(x["kod"]) in u:
            uy.append(uyari("Orta", "Mükerrer ürün kodu", x["kod"], "Ürün dosyasında birden çok kez geçiyor; ilk satır kullanıldı."))
            continue
        u[katla(x["kod"])] = {**x, "sayi": 0, "miktar": Z, "listede": True}
    for c in cikislar:
        e = u.setdefault(katla(c["kod"]), {"kod": c["kod"], "ad": "", "konum": "", "bolge": "", "sayi": 0, "miktar": Z, "listede": False})
        e["sayi"] += c["sayi"]
        e["miktar"] += c["miktar"]
    toplam = sum(e["sayi"] for e in u.values())
    if not toplam:
        raise ValueError("Çıkış kaydı yok (toplam çıkış sayısı sıfır).")
    liste = sorted(u.values(), key=lambda e: (-e["sayi"], -e["miktar"], katla(e["kod"])))
    kum = 0
    for e in liste:
        kim = f"{e['kod']} · {e['ad']}".strip(" ·")
        onceki = Decimal(kum * 100) / toplam
        kum += e["sayi"]
        e["pay"], e["kumulatif"] = yuvarla(Decimal(e["sayi"] * 100) / toplam), yuvarla(Decimal(kum * 100) / toplam)
        e["sinif"] = "D" if not e["sayi"] else "A" if onceki < a else "B" if onceki < b else "C"
        e["onerilen"] = "C" if e["sinif"] == "D" else e["sinif"]
        e["oneri"] = "Bölge bilinmiyor" if not e["bolge"] else "Yerinde" if e["bolge"] == e["onerilen"] else "Öne al" if RANK[e["bolge"]] > RANK[e["onerilen"]] else "Geriye al"
        if urunler is not None and not e["listede"]:
            uy.append(uyari("Bilgi", "Ürün listesinde yok", kim, f"Çıkış kaydında {e['sayi']} kez geçiyor ama ürün dosyasında yok; mevcut konumu bilinmiyor."))
        if e["sinif"] == "A" and e["bolge"] == "C":
            uy.append(uyari("Orta", "Sık çıkan ürün uzak bölgede", kim, f"Dönemde {e['sayi']} çıkış (toplamın %{e['pay']}'i) ama C bölgesinde ({e['konum'] or 'konum yazılmamış'}). Öne almak toplama yolunu kısaltır."))
        if e["sinif"] == "D" and e["bolge"] == "A":
            uy.append(uyari("Orta", "Çıkışı olmayan ürün ön bölgede", kim, f"Dönemde hiç çıkışı yok ama A bölgesinde ({e['konum'] or 'konum yazılmamış'}) yer tutuyor."))
    siniflar = [{"sinif": s, "urun": sum(1 for e in liste if e["sinif"] == s), "urun_pay": yuvarla(Decimal(sum(1 for e in liste if e["sinif"] == s) * 100) / len(liste)), "sayi": sum(e["sayi"] for e in liste if e["sinif"] == s),
                 "sayi_pay": yuvarla(Decimal(sum(e["sayi"] for e in liste if e["sinif"] == s) * 100) / toplam)} for s in "ABCD"]
    matris = [{"bolge": bl or "(bilinmiyor)", **{s: sum(1 for e in liste if e["bolge"] == bl and e["sinif"] == s) for s in "ABCD"}} for bl in ("A", "B", "C", "") if bl or any(not e["bolge"] for e in liste)]
    tasima = sorted((e for e in liste if e["oneri"] == "Öne al"), key=lambda e: (-e["sayi"], katla(e["kod"]))) + sorted((e for e in liste if e["oneri"] == "Geriye al"), key=lambda e: (e["sayi"], katla(e["kod"])))
    tarihler = [c["tarih"] for c in cikislar if c["tarih"]]
    return {"urunler": liste, "toplam": toplam, "siniflar": siniflar, "matris": matris, "tasima": tasima, "say": {o: sum(1 for e in liste if e["oneri"] == o) for o in ONERI_RENK}, "donem": (min(tarihler), max(tarihler)) if tarihler else None,
            "sinirlar": (a, b), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sınıflandırma"
    basliklar = ["Stok Kodu", "Stok Adı", "Çıkış Sayısı", "Çıkış Miktarı", "Pay %", "Kümülatif %", "Sınıf", "Mevcut Konum", "Mevcut Bölge", "Önerilen Bölge", "Öneri"]
    satir = lambda e: [e["kod"], e["ad"], e["sayi"], e["miktar"], e["pay"], e["kumulatif"], e["sinif"], e["konum"], e["bolge"], e["onerilen"], e["oneri"]]   # noqa: E731
    _baslik(ws, basliklar, (14, 30, 10, 13, 9, 11, 7, 16, 9, 10, 16))
    for e in r["urunler"]:
        _satir(ws, satir(e), para_s=(4, 5, 6), dolgu={7: {"A": YESIL, "B": SARI, "D": KIRMIZI}.get(e["sinif"]), 11: ONERI_RENK[e["oneri"]]})
    ws = wb.create_sheet("Taşıma Listesi")
    _baslik(ws, basliklar + ["Yeni Konum (siz doldurun)"], (14, 30, 10, 13, 9, 11, 7, 16, 9, 10, 16, 22))
    for e in r["tasima"]:
        _satir(ws, satir(e) + [None], para_s=(4, 5, 6), dolgu={11: ONERI_RENK[e["oneri"]]})
    ws = wb.create_sheet("Özet")
    _baslik(ws, ["Sınıf", "Ürün Sayısı", "Ürün Payı %", "Çıkış Sayısı", "Çıkış Payı %"], (36, 11, 11, 12, 12))
    for s in r["siniflar"]:
        _satir(ws, [s["sinif"], s["urun"], s["urun_pay"], s["sayi"], s["sayi_pay"]], para_s=(3, 5))
    ws.append([])
    _baslik(ws, ["Mevcut Bölge ↓ / Sınıf →", "A", "B", "C", "D"], dondur="A2")
    for m in r["matris"]:
        _satir(ws, [m["bolge"], m["A"], m["B"], m["C"], m["D"]])
    ws.append([])
    a, b = r["sinirlar"]
    ws.append([f"A: önceki kümülatif pay %{a:f} altında · B: %{b:f} altında · C: kalan · D: dönemde çıkışı yok. Toplam çıkış sayısı {r['toplam']}."
               + (f" Dönem: {r['donem'][0]:%d.%m.%Y} – {r['donem'][1]:%d.%m.%Y}." if r["donem"] else "")])
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(c_yol: Path, u_yol: Path | None, cikti: Path, a: Decimal = Decimal(80), b: Decimal = Decimal(95)) -> dict:
    cikislar, uy1 = oku_cikislar(c_yol)
    urunler, uy2 = oku_urunler(u_yol) if u_yol else (None, [])
    r = onerile(cikislar, urunler, a, b)
    r["uyarilar"] = sirala(uy1 + uy2 + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Ürünleri çıkış sıklığına göre ABC sınıflarına ayırıp yerleşim önerisi verir.")
    p.add_argument("--cikislar", type=Path, help="Çıkış satırları: Tarih, Stok Kodu, Miktar (her satır bir çıkış) ya da özet: Stok Kodu, Çıkış Sayısı, Miktar")
    p.add_argument("--urunler", type=Path, help="Stok Kodu, Stok Adı, Mevcut Konum, Bölge (A/B/C) (isteğe bağlı)")
    p.add_argument("--a", default="80", help="A sınıfı kümülatif pay sınırı, yüzde (varsayılan 80)")
    p.add_argument("--b", default="95", help="B sınıfı kümülatif pay sınırı, yüzde (varsayılan 95)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "depo_yerlesim_abc.xlsx")
    args = p.parse_args(argv)
    ornek = args.cikislar is None
    c_yol = ORNEK / "cikislar.csv" if ornek else args.cikislar
    u_yol = ORNEK / "urunler.csv" if ornek and args.urunler is None else args.urunler
    a, b = para(args.a), para(args.b)
    if a is None or b is None or not 0 < a < b <= 100:
        print("[X] --a ve --b yüzde olmalı; 0 < a < b ≤ 100")
        return 2
    for y in (c_yol, u_yol):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(c_yol, u_yol, args.cikti, a, b)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    s = {x["sinif"]: x for x in r["siniflar"]}
    print(f"[OK] {len(r['urunler'])} ürün · {r['toplam']} çıkış · A {s['A']['urun']} ürün (çıkışın %{s['A']['sayi_pay']}'i), B {s['B']['urun']}, C {s['C']['urun']}, D {s['D']['urun']}"
          + (f" · öne al {r['say']['Öne al']}, geriye al {r['say']['Geriye al']}" if u_yol else ""))
    print(f"[OK] Rapor: {args.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
