import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def c(kod, sayi=1, miktar=0, t=None):
    return {"tarih": t, "kod": kod, "sayi": sayi, "miktar": D(miktar)}


def u(kod, bolge="", konum="R1"):
    return {"kod": kod, "ad": "Ad", "konum": konum, "bolge": bolge}


class OneriTesti(unittest.TestCase):
    def test_sinif_sinirlari(self):
        l = [c("P4", 5), c("P2", 30), c("P1", 20), c("P3", 15), c("p1", 30, 7)]
        s = lambda **kw: [(e["kod"], e["sayi"], e["pay"], e["kumulatif"], e["sinif"]) for e in main.onerile(l, **kw)["urunler"]]   # noqa: E731
        self.assertEqual(s(), [("P1", 50, D("50.00"), D("50.00"), "A"), ("P2", 30, D("30.00"), D("80.00"), "A"), ("P3", 15, D("15.00"), D("95.00"), "B"), ("P4", 5, D("5.00"), D("100.00"), "C")])
        self.assertEqual(([x[4] for x in s(a=D(81))], [x[4] for x in s(b=D(96))], [x[4] for x in s(a=D(50), b=D(80))]), (["A", "A", "A", "C"], ["A", "A", "B", "B"], ["A", "B", "C", "C"]))
        r = main.onerile([c("b", 2, 5), c("a", 2, 5), c("z", 2, 9), c("tek", 1)])
        self.assertEqual(([e["kod"] for e in r["urunler"]], r["toplam"], r["urunler"][0]["oneri"], r["donem"], r["say"]["Bölge bilinmiyor"]), (["z", "a", "b", "tek"], 7, "Bölge bilinmiyor", None, 4))
        self.assertEqual([e["sayi"] for e in main.onerile([c("x"), c("x"), c("X"), c("y")])["urunler"]], [3, 1])
        r = main.onerile([c("x", t=date(2026, 9, 5)), c("x", t=date(2026, 8, 1)), c("x")])
        self.assertEqual(r["donem"], (date(2026, 8, 1), date(2026, 9, 5)))
        for bos in ([], [c("x", 0)]):
            with self.assertRaises(ValueError):
                main.onerile(bos, [u("x")])

    def test_yerlesim(self):
        cik = [c("A1", 40), c("A2", 30), c("A3", 10), c("B1", 10), c("B2", 5), c("C1", 3), c("C2", 2), c("YOK", 0)]
        ur = [u("A1", "A"), u("A2", "B"), u("A3", "C"), u("B1", "A"), u("B2", "B"), u("C1", "A"), u("c2", ""), u("D1", "A"), u("D2", "C"), u("D3", "B"), u("A1", "C")]
        r = main.onerile(cik + [c("EK", 0, 0)], ur)
        self.assertEqual([(e["kod"], e["sinif"], e["bolge"], e["onerilen"], e["oneri"]) for e in r["urunler"]],
                         [("A1", "A", "A", "A", "Yerinde"), ("A2", "A", "B", "A", "Öne al"), ("A3", "A", "C", "A", "Öne al"), ("B1", "B", "A", "B", "Geriye al"), ("B2", "B", "B", "B", "Yerinde"), ("C1", "C", "A", "C", "Geriye al"),
                          ("c2", "C", "", "C", "Bölge bilinmiyor"), ("D1", "D", "A", "C", "Geriye al"), ("D2", "D", "C", "C", "Yerinde"), ("D3", "D", "B", "C", "Geriye al"), ("EK", "D", "", "C", "Bölge bilinmiyor"),
                          ("YOK", "D", "", "C", "Bölge bilinmiyor")])
        self.assertEqual([e["kod"] for e in r["tasima"]], ["A2", "A3", "D1", "D3", "C1", "B1"])
        self.assertEqual(r["siniflar"], [{"sinif": "A", "urun": 3, "urun_pay": D("25.00"), "sayi": 80, "sayi_pay": D("80.00")}, {"sinif": "B", "urun": 2, "urun_pay": D("16.67"), "sayi": 15, "sayi_pay": D("15.00")},
                                         {"sinif": "C", "urun": 2, "urun_pay": D("16.67"), "sayi": 5, "sayi_pay": D("5.00")}, {"sinif": "D", "urun": 5, "urun_pay": D("41.67"), "sayi": 0, "sayi_pay": D("0.00")}])
        self.assertEqual(r["matris"], [{"bolge": "A", "A": 1, "B": 1, "C": 1, "D": 1}, {"bolge": "B", "A": 1, "B": 1, "C": 0, "D": 1}, {"bolge": "C", "A": 1, "B": 0, "C": 0, "D": 1}, {"bolge": "(bilinmiyor)", "A": 0, "B": 0, "C": 1, "D": 2}])
        self.assertEqual(r["say"], {"Öne al": 2, "Geriye al": 4, "Yerinde": 3, "Bölge bilinmiyor": 3})
        self.assertEqual(sorted((w["onem"], w["tur"], w["kim"]) for w in r["uyarilar"]),
                         sorted([("Orta", "Mükerrer ürün kodu", "A1"), ("Orta", "Sık çıkan ürün uzak bölgede", "A3 · Ad"), ("Orta", "Çıkışı olmayan ürün ön bölgede", "D1 · Ad"), ("Bilgi", "Ürün listesinde yok", "EK"),
                                 ("Bilgi", "Ürün listesinde yok", "YOK")]))
        r = main.onerile(cik)
        self.assertEqual((r["uyarilar"], [m["bolge"] for m in r["matris"]], r["tasima"]), ([], ["A", "B", "C", "(bilinmiyor)"], []))
        self.assertEqual([m["bolge"] for m in main.onerile([c("A1")], [u("A1", "A")])["matris"]], ["A", "B", "C"])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "cikislar.csv", ORNEK / "urunler.csv", cikti)
            self.assertEqual((r["toplam"], r["donem"]), (1000, (date(2026, 8, 31), date(2026, 9, 30))))
            self.assertEqual([(e["kod"], e["sayi"], e["pay"], e["kumulatif"], e["sinif"], e["bolge"], e["oneri"]) for e in r["urunler"]],
                             [("K-01", 320, D("32.00"), D("32.00"), "A", "A", "Yerinde"), ("K-02", 240, D("24.00"), D("56.00"), "A", "B", "Öne al"), ("K-03", 150, D("15.00"), D("71.00"), "A", "C", "Öne al"),
                              ("K-04", 90, D("9.00"), D("80.00"), "A", "A", "Yerinde"), ("K-05", 70, D("7.00"), D("87.00"), "B", "A", "Geriye al"), ("K-06", 50, D("5.00"), D("92.00"), "B", "B", "Yerinde"),
                              ("K-07", 30, D("3.00"), D("95.00"), "B", "C", "Öne al"), ("K-08", 20, D("2.00"), D("97.00"), "C", "A", "Geriye al"), ("K-09", 15, D("1.50"), D("98.50"), "C", "C", "Yerinde"),
                              ("K-10", 10, D("1.00"), D("99.50"), "C", "B", "Geriye al"), ("K-11", 5, D("0.50"), D("100.00"), "C", "C", "Yerinde"), ("K-12", 0, D("0.00"), D("100.00"), "D", "A", "Geriye al")])
            self.assertEqual((r["urunler"][0]["miktar"], r["urunler"][1]["miktar"]), (D(3840), D(1220)))
            self.assertEqual([(s["sinif"], s["urun"], s["urun_pay"], s["sayi"], s["sayi_pay"]) for s in r["siniflar"]],
                             [("A", 4, D("33.33"), 800, D("80.00")), ("B", 3, D("25.00"), 150, D("15.00")), ("C", 4, D("33.33"), 50, D("5.00")), ("D", 1, D("8.33"), 0, D("0.00"))])
            self.assertEqual(r["matris"], [{"bolge": "A", "A": 2, "B": 1, "C": 1, "D": 1}, {"bolge": "B", "A": 1, "B": 1, "C": 1, "D": 0}, {"bolge": "C", "A": 1, "B": 1, "C": 2, "D": 0}])
            self.assertEqual(([e["kod"] for e in r["tasima"]], r["say"]), (["K-02", "K-03", "K-07", "K-12", "K-10", "K-08", "K-05"], {"Öne al": 3, "Geriye al": 4, "Yerinde": 5, "Bölge bilinmiyor": 0}))
            self.assertEqual([(w["onem"], w["tur"], w["kim"]) for w in r["uyarilar"]],
                             [("Orta", "Sık çıkan ürün uzak bölgede", "K-03 · Kargo Poşeti Orta"), ("Orta", "Çıkışı olmayan ürün ön bölgede", "K-12 · Ahşap Sandık Seti")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Sınıflandırma", "Taşıma Listesi", "Özet", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            cy, uy_ = Path(tmp) / "c.csv", Path(tmp) / "u.csv"
            cy.write_text("Çıkış Tarihi;Ürün Kodu;Çıkış Miktarı\n05.09.2026;X1;2,5\n06.09.2026;X1;1.250,5\nx;X2;\n07.09.2026;X3;-1\n;;3\n", encoding="utf-8")
            uy_.write_text("Ürün Kodu;Ürün Adı;Raf;Zon\nX1;Bant;R-1;a\nX2;Film;;B Bölgesi\nX3;Koli;;Bölge C\nX4;Poşet;;Ön\n;Y;;A\n", encoding="utf-8")
            l, uy = main.oku_cikislar(cy)
            self.assertEqual((l, [(w["tur"], w["kim"]) for w in uy]),
                             ([{"tarih": date(2026, 9, 5), "kod": "X1", "sayi": 1, "miktar": D("2.5")}, {"tarih": date(2026, 9, 6), "kod": "X1", "sayi": 1, "miktar": D("1250.5")}, {"tarih": None, "kod": "X2", "sayi": 1, "miktar": D(0)}],
                              [("Çıkış satırı okunamadı", "X3")]))
            l, uy = main.oku_urunler(uy_)
            self.assertEqual(([(x["kod"], x["ad"], x["konum"], x["bolge"]) for x in l], [(w["tur"], w["kim"]) for w in uy]),
                             ([("X1", "Bant", "R-1", "A"), ("X2", "Film", "", "B"), ("X3", "Koli", "", "C"), ("X4", "Poşet", "", "")], [("Bölge tanımsız", "X4")]))
            cy2 = Path(tmp) / "c2.csv"
            cy2.write_text("Stok Kodu;Çıkış Sayısı\nX1;12\nX2;\nX3;1,5\nX4;-2\nX5;0\n", encoding="utf-8")
            l, uy = main.oku_cikislar(cy2)
            self.assertEqual(([(x["kod"], x["sayi"]) for x in l], [w["kim"] for w in uy]), ([("X1", 12), ("X2", 1), ("X5", 0)], ["X3", "X4"]))
            c_ = str(Path(tmp) / "x.xlsx")
            arg = ["--cikislar", str(cy), "--cikti", c_]
            self.assertEqual(main.main(arg + ["--urunler", str(uy_), "--a", "70", "--b", "90"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + x) for x in (["--a", "95", "--b", "80"], ["--a", "0"], ["--b", "101"], ["--a", "x"])], [2] * 4)
            self.assertEqual(main.main(arg + ["--urunler", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Tarih;Miktar\n05.09.2026;1\n", "Stok Kodu;Miktar\n;1\n"):
                cy.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c_]), 0)


if __name__ == "__main__":
    unittest.main()
