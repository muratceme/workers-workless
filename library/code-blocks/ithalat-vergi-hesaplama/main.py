"""
İthalat Vergi Hesaplama — Workers / Workless kod bloğu
Gümrük Operasyon › Gümrük Operasyon Uzmanı

CIF kıymet, gümrük vergisi, ilave gümrük vergisi ve KDV oranlarından ithalatta ödenecek vergileri hesaplar.
  CIF (TL)        = CIF kıymet (döviz) × --kur                                       (kıymet zaten TL ise kur 1)
  Gümrük vergisi  = CIF × GV %;      ilave gümrük vergisi (İGV) = CIF × İGV %
  KKDF            = CIF × --kkdf %                                                   (yalnız vadeli ödeme şekillerinde; varsayılan 0)
  ÖTV matrahı     = CIF + GV + İGV + KKDF + diğer vergi + damga vergisi payı + yurt içi gider payı;   ÖTV = matrah × ÖTV %   (yalnız oransal ÖTV)
  KDV matrahı     = ÖTV matrahı + ÖTV;                                               KDV = matrah × KDV %
  Paylar          = --damga ve --yurt-ici-gider tutarları kalemlere CIF oranında dağıtılır (kuruş farkı son kalemde)
  Toplam vergi    = GV + İGV + KKDF + diğer vergi + damga + ÖTV + KDV;   KDV hariç maliyet = CIF + yurt içi gider + (toplam vergi − KDV)
Oranları siz verirsiniz: kod GTİP'e göre vergi oranı, menşe, tercihli rejim, muafiyet veya gözetim uygulaması bilmez. Maktu ÖTV (litre, adet başına) hesaplanmaz.
Sonuç ön hesaptır; beyanname gümrük müşavirince hazırlanır ve gümrük idaresinin tahakkuku esastır.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"kalem": ("kalem", "gtip", "esya", "urun", "esya tanimi"), "aciklama": ("aciklama", "tanim", "esya adi"), "cif": ("cif kiymet", "cif", "kiymet", "gumruk kiymeti"),
         "gv": ("gv", "gumruk vergisi", "gv orani", "gumruk vergisi orani"), "igv": ("igv", "ilave gumruk vergisi", "igv orani"), "otv": ("otv", "otv orani"), "kdv": ("kdv", "kdv orani"),
         "diger": ("diger vergi", "dampin", "anti damping", "diger vergi tutari")}
VERGI = (("gv", "Gümrük Vergisi"), ("igv", "İlave Gümrük Vergisi"), ("kkdf", "KKDF"), ("diger", "Diğer Vergi"), ("damga", "Damga Vergisi"), ("otv", "ÖTV"), ("kdv", "KDV"))
Z = Decimal(0)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("cif", "kdv")):
        if not any(metin(r[a]) for a in ("kalem", "cif")):
            continue
        kim = f"{metin(r['kalem']) or metin(r['aciklama']) or 'kalem'} (satır {r['_satir']})"
        cif = para(r["cif"])
        if cif is None or cif <= 0:
            uy.append(uyari("Yüksek", "CIF kıymet okunamadı", kim, f"'{metin(r['cif'])}'. Kalem alınmadı."))
            continue
        x, hata = {"kalem": metin(r["kalem"]), "aciklama": metin(r["aciklama"]), "cif_doviz": cif}, False
        for a in ("gv", "igv", "otv", "kdv"):
            v = para(r[a]) if metin(r[a]) else (None if a == "kdv" else Z)
            if v is None or not 0 <= v <= 1000:
                uy.append(uyari("Yüksek", "Oran okunamadı", kim, f"{a.upper()} oranı '{metin(r[a])}'. Kalem alınmadı" + (": KDV oranı boş bırakılamaz (istisna ise 0 yazın)." if a == "kdv" else ".")))
                hata = True
                break
            x[a + "_oran"] = v
        if hata:
            continue
        d = para(r["diger"])
        if d is None and metin(r["diger"]) or (d is not None and d < 0):
            uy.append(uyari("Yüksek", "Diğer vergi okunamadı", kim, f"'{metin(r['diger'])}'. 0 sayıldı."))
            d = None
        x["diger"] = d or Z
        l.append(x)
    return l, uy


def dagit(tutar: Decimal, agirliklar: list[Decimal]) -> list[Decimal]:
    toplam, kalan, sonuc = sum(agirliklar, Z), tutar, []
    for i, a in enumerate(agirliklar):
        p = kalan if i == len(agirliklar) - 1 else yuvarla(tutar * a / toplam)
        kalan -= p
        sonuc.append(p)
    return sonuc


def hesapla(kalemler: list[dict], kur: Decimal = Decimal(1), kkdf: Decimal = Z, damga: Decimal = Z, yurt_ici: Decimal = Z) -> dict:
    if not kalemler:
        raise ValueError("Hesaplanacak kalem yok.")
    for k in kalemler:
        k["cif"] = yuvarla(k["cif_doviz"] * kur)
    cifler = [k["cif"] for k in kalemler]
    for k, dp, yp in zip(kalemler, dagit(damga, cifler), dagit(yurt_ici, cifler)):
        c = k["cif"]
        k["gv"], k["igv"], k["kkdf"], k["damga"], k["yurt_ici"] = yuvarla(c * k["gv_oran"] / 100), yuvarla(c * k["igv_oran"] / 100), yuvarla(c * kkdf / 100), dp, yp
        k["otv_matrah"] = c + k["gv"] + k["igv"] + k["kkdf"] + k["diger"] + dp + yp
        k["otv"] = yuvarla(k["otv_matrah"] * k["otv_oran"] / 100)
        k["kdv_matrah"] = k["otv_matrah"] + k["otv"]
        k["kdv"] = yuvarla(k["kdv_matrah"] * k["kdv_oran"] / 100)
        k["toplam_vergi"] = sum((k[a] for a, _ in VERGI), Z)
        k["maliyet"] = c + yp + k["toplam_vergi"] - k["kdv"]
        k["yuk_orani"] = yuvarla((k["toplam_vergi"] - k["kdv"]) / c * 100)
    t = {a: sum((k[a] for k in kalemler), Z) for a in ("cif", "yurt_ici", "toplam_vergi", "maliyet", "kdv_matrah", *[v for v, _ in VERGI])}
    t["odenecek"] = t["toplam_vergi"]
    return {"kalemler": kalemler, "toplam": t, "ayar": (kur, kkdf, damga, yurt_ici)}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    kur, kkdf, damga, yi = r["ayar"]
    t = r["toplam"]
    ws = wb.active
    ws.title = "Vergi Özeti"
    _baslik(ws, ["Kalem", "Tutar (TL)", "Açıklama"], (30, 18, 80))
    for ad, d, ac in (("CIF kıymet", t["cif"], f"Kur {kur:f}"), *[(ad, t[a], "") for a, ad in VERGI if t[a] or a in ("gv", "kdv")], ("TOPLAM VERGİ", t["toplam_vergi"], "Gümrükte ödenecek vergiler toplamı"),
                      ("Yurt içi giderler", t["yurt_ici"], "KDV matrahına dahil edilen, sizin verdiğiniz tutar"), ("KDV matrahı", t["kdv_matrah"], ""),
                      ("KDV hariç maliyet", t["maliyet"], "CIF + yurt içi gider + KDV dışındaki vergiler (KDV indirim konusu yapılabiliyorsa)")):
        n = _satir(ws, [ad, d, ac], para_s=(2,), sar=(3,))
        if ad.isupper():
            for c in ws[n]:
                c.font = Font(bold=True)
            ws.cell(n, 2).fill = PatternFill("solid", fgColor=YESIL)
    ws.append([])
    ws.append(["Ön hesaptır. Oranlar sizin verdiğiniz değerlerdir; GTİP, menşe, tercihli rejim, muafiyet ve gözetim kontrolü yapılmaz. Gümrük idaresinin tahakkuku esastır."])
    ws = wb.create_sheet("Kalem Bazında")
    _baslik(ws, ["Kalem / GTİP", "Açıklama", "CIF (döviz)", "CIF (TL)", "GV %", "İGV %", "ÖTV %", "KDV %", *[ad for _, ad in VERGI[:5]], "Yurt İçi Gider Payı", "ÖTV Matrahı", "ÖTV", "KDV Matrahı", "KDV",
                 "Toplam Vergi", "KDV Hariç Maliyet", "KDV Hariç Vergi Yükü %"], (20, 30, 14, 15, 7, 7, 7, 7, 14, 14, 12, 12, 12, 12, 15, 13, 15, 13, 15, 15, 11))
    for k in r["kalemler"]:
        _satir(ws, [k["kalem"], k["aciklama"], k["cif_doviz"], k["cif"], k["gv_oran"], k["igv_oran"], k["otv_oran"], k["kdv_oran"], k["gv"], k["igv"], k["kkdf"], k["diger"], k["damga"], k["yurt_ici"],
                    k["otv_matrah"], k["otv"], k["kdv_matrah"], k["kdv"], k["toplam_vergi"], k["maliyet"], k["yuk_orani"]], para_s=(3, 4, *range(9, 22)))
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, **ayar) -> dict:
    kalemler, uy = oku(yol)
    r = hesapla(kalemler, **ayar)
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="İthalatta gümrük vergisi, İGV, ÖTV ve KDV ön hesabı yapar.")
    p.add_argument("--kalemler", type=Path, default=ORNEK / "ithalat_kalemleri.csv", help="Kalem / GTİP, Açıklama, CIF Kıymet, GV %%, İGV %%, ÖTV %%, KDV %%, Diğer Vergi (tutar)")
    p.add_argument("--kur", help="CIF kıymetin TL'ye çevrileceği kur (varsayılan 1; örnek veride 40)")
    p.add_argument("--kkdf", default="0", help="KKDF oranı %% (vadeli ödemede; varsayılan 0)")
    p.add_argument("--damga", default="0", help="Damga vergisi tutarı, TL (varsayılan 0)")
    p.add_argument("--yurt-ici-gider", help="KDV matrahına giren yurt içi giderler toplamı, TL (varsayılan 0; örnek veride 6.000)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "ithalat_vergileri.xlsx")
    a = p.parse_args(argv)
    ornek = a.kalemler == ORNEK / "ithalat_kalemleri.csv"
    d = [para(x) for x in (a.kur or ("40" if ornek else "1"), a.kkdf, a.damga, a.yurt_ici_gider or ("6000" if ornek else "0"))]
    if any(x is None or x < 0 for x in d) or d[0] == 0 or d[1] > 100:
        print("[X] --kur pozitif; --kkdf 0-100 arasında; --damga ve --yurt-ici-gider sıfır veya pozitif sayı olmalı")
        return 2
    if not a.kalemler.exists():
        print(f"[X] Dosya bulunamadı: {a.kalemler}")
        return 1
    try:
        r = calistir(a.kalemler, a.cikti, kur=d[0], kkdf=d[1], damga=d[2], yurt_ici=d[3])
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {len(r['kalemler'])} kalem · CIF {t['cif']:,.2f} TL · " + " · ".join(f"{ad} {t[v]:,.2f}" for v, ad in VERGI if t[v]))
    print(f"[OK] Toplam vergi {t['toplam_vergi']:,.2f} TL · KDV hariç maliyet {t['maliyet']:,.2f} TL (ön hesap)")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
