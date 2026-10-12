import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def kl(cif, gv=0, igv=0, otv=0, kdv=20, diger=0):
    return {"kalem": "K", "aciklama": "", "cif_doviz": D(cif), "gv_oran": D(gv), "igv_oran": D(igv), "otv_oran": D(otv), "kdv_oran": D(kdv), "diger": D(diger)}


def tek(*a, **kw):
    ayar = {k: D(kw.pop(k)) for k in ("kur", "kkdf", "damga", "yurt_ici") if k in kw}
    return main.hesapla([kl(*a, **kw)], **ayar)["kalemler"][0]


class HesapTesti(unittest.TestCase):
    def test_vergiler(self):
        k = tek(10000, 5, 20, kur=40, yurt_ici=4000)
        self.assertEqual((k["cif"], k["gv"], k["igv"], k["otv"], k["kdv_matrah"], k["kdv"], k["toplam_vergi"], k["maliyet"], k["yuk_orani"]),
                         (D(400000), D(20000), D(80000), D(0), D(504000), D(100800), D(200800), D(504000), D("25.00")))
        k = tek(5000, otv=20, kur=40, yurt_ici=2000)
        self.assertEqual((k["otv_matrah"], k["otv"], k["kdv_matrah"], k["kdv"], k["toplam_vergi"]), (D(202000), D(40400), D(242400), D(48480), D(88880)))
        k = tek(1000, 10, kkdf=6, damga=50, diger=30)
        self.assertEqual((k["gv"], k["kkdf"], k["kdv_matrah"], k["kdv"], k["toplam_vergi"]), (D(100), D(60), D(1240), D(248), D(488)))      # 1000 + 100 + 60 + 30 + 50
        k = tek(1000, kdv=0)
        self.assertEqual((k["kdv"], k["toplam_vergi"], k["maliyet"]), (D(0), D(0), D(1000)))
        self.assertEqual(tek("333.33", "2.5", kdv=10)["gv"], D("8.33"))

    def test_dagitim_ve_toplam(self):
        self.assertEqual((main.dagit(D(6000), [D(400000), D(200000)]), main.dagit(D(100), [D(1), D(1), D(1)]), main.dagit(D(0), [D(5)])), ([D(4000), D(2000)], [D("33.33"), D("33.33"), D("33.34")], [D(0)]))
        r = main.hesapla([kl(10000, 5, 20), kl(5000, otv=20)], D(40), yurt_ici=D(6000))
        t = r["toplam"]
        self.assertEqual((t["cif"], t["gv"], t["igv"], t["otv"], t["kdv"], t["toplam_vergi"], t["maliyet"], t["kdv_matrah"]), (D(600000), D(20000), D(80000), D(40400), D(149280), D(289680), D(746400), D(746400)))
        with self.assertRaises(ValueError):
            main.hesapla([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "i.xlsx"
            r = main.calistir(ORNEK / "ithalat_kalemleri.csv", cikti, kur=D(40), yurt_ici=D(6000))
            self.assertEqual((r["toplam"]["toplam_vergi"], r["toplam"]["maliyet"], [k["toplam_vergi"] for k in r["kalemler"]], r["uyarilar"]), (D(289680), D(746400), [D(200800), D(88880)], []))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Vergi Özeti", "Kalem Bazında", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "k.csv"
            y.write_text("GTİP;CIF;GV;KDV;Dampin\nA;1.250,50;%5;20;\nB;yok;0;20;\nC;100;;;\nD;100;çok;20;\nE;100;;0;az\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual([(k["kalem"], k["cif_doviz"], k["gv_oran"], k["igv_oran"], k["kdv_oran"], k["diger"]) for k in l], [("A", D("1250.50"), D(5), D(0), D(20), D(0)), ("E", D(100), D(0), D(0), D(0), D(0))])
            self.assertEqual(sorted(u["tur"] for u in uy), ["CIF kıymet okunamadı", "Diğer vergi okunamadı", "Oran okunamadı", "Oran okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--kalemler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--kur", "41,25", "--kkdf", "6", "--damga", "1.200,50", "--yurt-ici-gider", "5.000"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(arg + ["--kur", "0"]), 2)
            self.assertEqual(main.main(arg + ["--kkdf", "150"]), 2)
            self.assertEqual(main.main(["--kalemler", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("CIF;KDV\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
