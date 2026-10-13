import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
DONEM, BUGUN = (2026, 9), date(2026, 10, 13)


def b(mukellef, belge, sorumlu="S1"):
    return {"mukellef": mukellef, "belge": belge, "sorumlu": sorumlu, "iletisim": ""}


def g(mukellef, belge, gun=None, donem=DONEM, ay=10):
    return {"mukellef": mukellef, "belge": belge, "donem": donem, "tarih": date(2026, ay, gun) if gun else None, "not": ""}


BEK = lambda: [b("A", "Satış"), b("A", "Alış"), b("A", "Banka"), b("A", "Puantaj"), b("B", "Satış", "S2"), b("B", "Alış", "S2"), b("C", "Satış", ""), b("C", "Kasa", "")]   # noqa: E731
GEL = lambda: [g("A", "Satış", 5), g("a", "ALIŞ", 12), g("A", "Banka", 4, (2026, 8), 9), g("B", "Satış", 2), g("B", "Alış"), g("D", "Satış", 3)]   # noqa: E731


class TakipTesti(unittest.TestCase):
    def test_donem_cozme(self):
        self.assertEqual([main.donem_coz(x) for x in ("2026-09", "2026/9", "09.2026", "9/2026", "Eylül 2026", "EYLUL 2026", "01.09.2026", "2026-09-30", date(2026, 9, 1), "2026-13", "Eylül", "", "26-09")],
                         [(2026, 9)] * 9 + [None] * 4)

    def test_mukellef_durumu(self):
        r = main.takip(BEK(), GEL(), DONEM, BUGUN)
        self.assertEqual([(m["mukellef"], m["durum"], m["beklenen"], m["gelen"], m["gec"], m["tamamlanma"], m["eksik"]) for m in r["mukellefler"]],
                         [("C", "Hiç gelmedi", 2, 0, 0, D("0.00"), ["Satış", "Kasa"]), ("A", "Eksik var", 4, 2, 1, D("50.00"), ["Banka", "Puantaj"]), ("B", "Tamam", 2, 2, 0, D("100.00"), [])])
        self.assertEqual([(s["mukellef"], s["belge"], s["durum"], s["gecikme"]) for s in r["satirlar"]],
                         [("A", "Satış", "Geldi", 0), ("A", "Alış", "Geç geldi", 0), ("A", "Banka", "Gecikmiş", 3), ("A", "Puantaj", "Gecikmiş", 3), ("B", "Satış", "Geldi", 0), ("B", "Alış", "Geldi", 0), ("C", "Satış", "Gecikmiş", 3),
                          ("C", "Kasa", "Gecikmiş", 3)])
        o = r["ozet"]
        self.assertEqual((o["donem"], o["son"], o["gecikme"], o["mukellef"], o["beklenen"], o["gelen"], o["eksik"], o["tamamlanma"], o["say"], r["turler"]),
                         ("Eylül 2026", date(2026, 10, 10), 3, 3, 8, 4, 4, D("50.00"), {"Tamam": 1, "Eksik var": 1, "Hiç gelmedi": 1}, ["Satış", "Alış", "Banka", "Puantaj", "Kasa"]))
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]), [("Bilgi", "Beklenen listesinde olmayan belge", "D · Satış"), ("Orta", "Gecikmiş belge", "A"), ("Yüksek", "Gecikmiş belge", "C")])
        self.assertEqual(([(s["sorumlu"], s["beklenen"], s["eksik"], s["mukellefler"]) for s in r["sorumlular"]], r["mukellefler"][2]["hatirlatma"]), ([("Atanmamış", 2, 2, ["C"]), ("S1", 4, 2, ["A"]), ("S2", 2, 0, [])], ""))
        self.assertEqual(r["mukellefler"][1]["hatirlatma"], "Sayın A, Eylül 2026 dönemine ait şu belgeleriniz henüz bize ulaşmadı: Banka, Puantaj. Kayıtlarınızın zamanında işlenebilmesi için belgeleri en kısa sürede iletmenizi rica ederiz.")

    def test_tarih_ve_son_gun(self):
        r = main.takip(BEK(), GEL(), DONEM, date(2026, 10, 10))
        self.assertEqual((r["ozet"]["gecikme"], {s["durum"] for s in r["satirlar"]}, r["uyarilar"][1:]), (0, {"Geldi", "Geç geldi", "Bekleniyor"}, []))
        r = main.takip(BEK(), GEL(), DONEM, BUGUN, son_gun=12)
        self.assertEqual((r["ozet"]["son"], r["ozet"]["gecikme"], r["satirlar"][1]["durum"], r["mukellefler"][1]["gec"]), (date(2026, 10, 12), 1, "Geldi", 0))
        r = main.takip(BEK(), GEL(), (2026, 12), date(2027, 1, 5))
        self.assertEqual((r["ozet"]["donem"], r["ozet"]["son"], r["ozet"]["gelen"], r["ozet"]["say"]["Hiç gelmedi"], r["satirlar"][0]["durum"]), ("Aralık 2026", date(2027, 1, 10), 0, 3, "Bekleniyor"))
        r = main.takip(BEK(), [g("A", "Satış", 20), g("A", "Satış", 8)], DONEM, BUGUN)
        self.assertEqual((r["satirlar"][0]["durum"], r["satirlar"][0]["tarih"]), ("Geldi", date(2026, 10, 8)))                 # aynı belge iki kez geldiyse ilk geliş esas
        with self.assertRaises(ValueError):
            main.takip([], GEL(), DONEM, BUGUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "e.xlsx"
            r = main.calistir(ORNEK / "beklenen_belgeler.csv", ORNEK / "gelen_belgeler.csv", cikti, DONEM, BUGUN)
            self.assertEqual([(m["mukellef"], m["durum"], m["tamamlanma"], m["eksik"]) for m in r["mukellefler"]],
                             [("Numune Gıda (Şahıs)", "Hiç gelmedi", D("0.00"), ["Satış Faturaları", "Alış Faturaları", "Kasa Defteri"]), ("Örnek Ticaret Ltd. Şti.", "Eksik var", D("50.00"), ["Banka Ekstresi", "Puantaj"]),
                              ("Deneme İnşaat A.Ş.", "Tamam", D("100.00"), [])])
            self.assertEqual((r["ozet"]["beklenen"], r["ozet"]["gelen"], r["ozet"]["tamamlanma"], [(u["tur"], u["kim"]) for u in r["uyarilar"]]),
                             (10, 5, D("50.00"), [("Gecikmiş belge", "Numune Gıda (Şahıs)"), ("Gecikmiş belge", "Örnek Ticaret Ltd. Şti."), ("Beklenen listesinde olmayan belge", "Deneme İnşaat A.Ş. · Puantaj")]))
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Mükellef Özeti", "Belge Matrisi", "Eksik Belgeler", "Sorumlu Özeti", "Uyarılar"])
            self.assertEqual([c.value for c in wb["Belge Matrisi"][2]], ["Deneme İnşaat A.Ş.", "Geldi", "Geldi", "Geldi", "—", "—"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            by, gy = Path(tmp) / "b.csv", Path(tmp) / "g.csv"
            by.write_text("Firma;Evrak\nA;Satış\nA;satış\nA;\n;;\nB;Alış\n", encoding="utf-8")
            gy.write_text("Firma;Evrak;Ay;Tarih\nA;Satış;09.2026;05.10.2026\nB;Alış;geçen ay;\nB;;2026-09;\n", encoding="utf-8")
            bek, gel, uy = main.oku(by, gy)
            self.assertEqual(([(x["mukellef"], x["belge"]) for x in bek], [(x["mukellef"], x["donem"], x["tarih"]) for x in gel], sorted(u["tur"] for u in uy)),
                             ([("A", "Satış"), ("B", "Alış")], [("A", (2026, 9), date(2026, 10, 5))], ["Beklenen belge satırı okunamadı", "Gelen belge satırı okunamadı", "Gelen belge satırı okunamadı", "Mükerrer beklenen belge"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--beklenen", str(by), "--cikti", c]
            self.assertEqual(main.main(arg + ["--gelen", str(gy), "--donem", "2026-09", "--tarih", "13.10.2026", "--son-gun", "15"]), 0)
            self.assertEqual(main.main(arg + ["--donem", "Eylül 2026"]), 0)
            self.assertEqual(main.main(arg), 2)
            for ek in (["--donem", "eylül"], ["--donem", "2026-09", "--son-gun", "31"], ["--donem", "2026-09", "--tarih", "x"]):
                self.assertEqual(main.main(arg + ek), 2)
            self.assertEqual(main.main(arg + ["--donem", "2026-09", "--gelen", str(Path(tmp) / "yok.csv")]), 1)
            by.write_text("Firma\nA\n", encoding="utf-8")
            self.assertEqual(main.main(arg + ["--donem", "2026-09"]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
