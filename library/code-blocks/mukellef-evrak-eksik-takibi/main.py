"""
Mükellef Evrak Eksik Takibi — Workers / Workless kod bloğu
Muhasebe ve Defter İşleri › Mali Müşavir Stajyeri

Her mükelleften dönem için beklenen belgeleri gelenlerle karşılaştırır; eksik listesini, tamamlanma oranını ve hatırlatma metnini çıkarır.
  Beklenen           = "beklenen belgeler" dosyasındaki her mükellef – belge türü satırı                 (her ay aynı liste kullanılır)
  Geldi              = gelen belge kaydında aynı mükellef, aynı belge türü ve seçilen dönem (--donem) varsa
  Son teslim tarihi  = dönemi izleyen ayın --son-gun günü (varsayılan 10)                                (büronuzun kendi iç kuralı; yasal süre değildir)
  Eksik durumu       = bugün ≤ son teslim → "Bekleniyor";  bugün > son teslim → "Gecikmiş" (gecikme günüyle)
  Geç geldi          = geliş tarihi son teslim tarihinden sonra
  Tamamlanma %       = gelen belge türü sayısı ÷ beklenen belge türü sayısı × 100
  Mükellef durumu    = Tamam · Eksik var · Hiç gelmedi
Kod belgenin içeriğini denetlemez; yalnız "geldi" kaydına bakar. Beyanname ve bildirim süreleri bu paketin konusu değildir.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


BEKLENEN = {"mukellef": ("mukellef", "musteri", "firma", "unvan", "mukellef adi"), "belge": ("belge turu", "belge", "evrak", "evrak turu"), "sorumlu": ("sorumlu", "sorumlu personel", "takip eden"),
            "iletisim": ("iletisim", "e posta", "eposta", "telefon")}
GELEN = {"mukellef": BEKLENEN["mukellef"], "belge": BEKLENEN["belge"], "donem": ("donem", "ay", "ait oldugu donem"), "tarih": ("gelis tarihi", "tarih", "teslim tarihi"), "not": ("not", "aciklama")}
AYLAR = ("Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık")
AY_NO = {katla(a): i for i, a in enumerate(AYLAR, 1)}
DURUM_RENK = {"Tamam": YESIL, "Eksik var": SARI, "Hiç gelmedi": KIRMIZI}


def donem_coz(x) -> tuple[int, int] | None:
    """2026-09 · 09.2026 · 09/2026 · Eylül 2026 · tarih hücresi → (yıl, ay)"""
    if isinstance(x, (date, datetime)):
        return x.year, x.month
    s = metin(x)
    m = re.fullmatch(r"(\d{4})[-./](\d{1,2})", s) or re.fullmatch(r"(\d{4})[-./](\d{1,2})[-./]\d{1,2}", s)
    if m:
        y, a = int(m.group(1)), int(m.group(2))
    else:
        m = re.fullmatch(r"(?:\d{1,2}[-./])?(\d{1,2})[-./](\d{4})", s)
        if m:
            y, a = int(m.group(2)), int(m.group(1))
        else:
            p = katla(s).split()
            if len(p) != 2 or p[0] not in AY_NO or not p[1].isdigit():
                return None
            y, a = int(p[1]), AY_NO[p[0]]
    return (y, a) if 1 <= a <= 12 and 2000 <= y <= 2100 else None


def oku(beklenen_yol: Path, gelen_yol: Path | None = None) -> tuple[list[dict], list[dict], list[dict]]:
    uy, beklenen, gelen, gorulen = [], [], [], set()
    for r in kayitlar(beklenen_yol, BEKLENEN, ("mukellef", "belge")):
        m, b = metin(r["mukellef"]), metin(r["belge"])
        if not m and not b:
            continue
        if not m or not b:
            uy.append(uyari("Yüksek", "Beklenen belge satırı okunamadı", f"{m or '?'} · {b or '?'} (satır {r['_satir']})", "Mükellef veya belge türü boş. Satır alınmadı."))
            continue
        if (katla(m), katla(b)) in gorulen:
            uy.append(uyari("Bilgi", "Mükerrer beklenen belge", f"{m} · {b}", f"Aynı mükellef ve belge türü birden fazla satırda (satır {r['_satir']}). Bir kez sayıldı."))
            continue
        gorulen.add((katla(m), katla(b)))
        beklenen.append({"mukellef": m, "belge": b, "sorumlu": metin(r["sorumlu"]), "iletisim": metin(r["iletisim"])})
    if gelen_yol:
        for r in kayitlar(gelen_yol, GELEN, ("mukellef", "belge", "donem")):
            m, b = metin(r["mukellef"]), metin(r["belge"])
            if not m and not b:
                continue
            d = donem_coz(r["donem"])
            if not m or not b or d is None:
                uy.append(uyari("Yüksek", "Gelen belge satırı okunamadı", f"{m or '?'} · {b or '?'} (satır {r['_satir']})", f"Dönem '{metin(r['donem'])}' okunamadı (ör. 2026-09, 09.2026, Eylül 2026) veya mükellef, belge türü boş. Satır alınmadı; belge eksik görünebilir."))
                continue
            gelen.append({"mukellef": m, "belge": b, "donem": d, "tarih": tarih(r["tarih"]), "not": metin(r["not"])})
    return beklenen, gelen, uy


def takip(beklenen: list[dict], gelen: list[dict], donem: tuple[int, int], bugun: date, son_gun: int = 10) -> dict:
    if not beklenen:
        raise ValueError("Beklenen belge listesi boş.")
    y, a = donem
    son = date(y + (a == 12), a % 12 + 1, son_gun)
    donem_adi, uy = f"{AYLAR[a - 1]} {y}", []
    gecikme = max((bugun - son).days, 0)
    geldi = {}
    for g in gelen:
        if g["donem"] != donem:
            continue
        k = (katla(g["mukellef"]), katla(g["belge"]))
        if k not in geldi or (g["tarih"] and (geldi[k]["tarih"] is None or g["tarih"] < geldi[k]["tarih"])):
            geldi[k] = g
    bek_anahtar = {(katla(b["mukellef"]), katla(b["belge"])) for b in beklenen}
    for k, g in geldi.items():
        if k not in bek_anahtar:
            uy.append(uyari("Bilgi", "Beklenen listesinde olmayan belge", f"{g['mukellef']} · {g['belge']}", "Belge gelmiş ama beklenen belgeler listesinde yok. Mükellef veya belge adı farklı yazılmış olabilir; listeyi güncelleyin."))
    mus, satirlar, turler = {}, [], []
    for b in beklenen:
        k = (katla(b["mukellef"]), katla(b["belge"]))
        if b["belge"] not in turler:
            turler.append(b["belge"])
        m = mus.setdefault(k[0], {"mukellef": b["mukellef"], "sorumlu": b["sorumlu"], "iletisim": b["iletisim"], "beklenen": 0, "gelen": 0, "gec": 0, "eksik": [], "belgeler": {}})
        m["sorumlu"], m["iletisim"] = m["sorumlu"] or b["sorumlu"], m["iletisim"] or b["iletisim"]
        m["beklenen"] += 1
        g = geldi.get(k)
        if g:
            gec = bool(g["tarih"] and g["tarih"] > son)
            m["gelen"] += 1
            m["gec"] += gec
            durum = "Geç geldi" if gec else "Geldi"
        else:
            durum = "Gecikmiş" if gecikme else "Bekleniyor"
            m["eksik"].append(b["belge"])
        m["belgeler"][katla(b["belge"])] = durum
        satirlar.append({"mukellef": b["mukellef"], "belge": b["belge"], "sorumlu": b["sorumlu"], "durum": durum, "tarih": g["tarih"] if g else None, "gecikme": gecikme if durum == "Gecikmiş" else 0, "not": g["not"] if g else ""})
    liste = []
    for m in mus.values():
        m["tamamlanma"] = yuvarla(Decimal(m["gelen"]) / m["beklenen"] * 100)
        m["durum"] = "Tamam" if not m["eksik"] else "Hiç gelmedi" if m["gelen"] == 0 else "Eksik var"
        m["hatirlatma"] = (f"Sayın {m['mukellef']}, {donem_adi} dönemine ait şu belgeleriniz henüz bize ulaşmadı: {', '.join(m['eksik'])}. "
                           f"Kayıtlarınızın zamanında işlenebilmesi için belgeleri en kısa sürede iletmenizi rica ederiz.") if m["eksik"] else ""
        if m["eksik"] and gecikme:
            uy.append(uyari("Yüksek" if m["gelen"] == 0 else "Orta", "Gecikmiş belge", m["mukellef"], f"{len(m['eksik'])} belge {gecikme} gündür gecikmede: {', '.join(m['eksik'])}." + (f" Sorumlu: {m['sorumlu']}." if m["sorumlu"] else "")))
        liste.append(m)
    liste.sort(key=lambda m: (m["tamamlanma"], m["mukellef"]))
    sorumlu = {}
    for s in satirlar:
        o = sorumlu.setdefault(katla(s["sorumlu"]), {"sorumlu": s["sorumlu"] or "Atanmamış", "beklenen": 0, "eksik": 0, "mukellefler": set()})
        o["beklenen"] += 1
        if s["durum"] in ("Gecikmiş", "Bekleniyor"):
            o["eksik"] += 1
            o["mukellefler"].add(s["mukellef"])
    top_b, top_g = sum(m["beklenen"] for m in liste), sum(m["gelen"] for m in liste)
    ozet = {"donem": donem_adi, "son": son, "bugun": bugun, "gecikme": gecikme, "mukellef": len(liste), "beklenen": top_b, "gelen": top_g, "eksik": top_b - top_g, "tamamlanma": yuvarla(Decimal(top_g) / top_b * 100),
            "say": {d: sum(1 for m in liste if m["durum"] == d) for d in ("Tamam", "Eksik var", "Hiç gelmedi")}}
    return {"mukellefler": liste, "satirlar": satirlar, "turler": turler, "sorumlular": sorted(({**o, "mukellefler": sorted(o["mukellefler"])} for o in sorumlu.values()), key=lambda o: (-o["eksik"], o["sorumlu"])),
            "ozet": ozet, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb, o = Workbook(), r["ozet"]
    ws = wb.active
    ws.title = "Mükellef Özeti"
    _baslik(ws, ["Mükellef", "Durum", "Beklenen", "Gelen", "Eksik", "Tamamlanma %", "Geç Gelen", "Eksik Belgeler", "Sorumlu", "İletişim", "Hatırlatma Metni", "Arandı / Not (siz doldurun)"], (32, 13, 9, 8, 8, 11, 9, 44, 18, 24, 80, 26))
    for m in r["mukellefler"]:
        _satir(ws, [m["mukellef"], m["durum"], m["beklenen"], m["gelen"], len(m["eksik"]) or None, m["tamamlanma"], m["gec"] or None, ", ".join(m["eksik"]), m["sorumlu"], m["iletisim"], m["hatirlatma"], ""], para_s=(6,),
               dolgu={2: DURUM_RENK[m["durum"]]}, sar=(8, 11))
    ws.append([])
    ws.append([f"{o['donem']} dönemi · son teslim {o['son']:%d.%m.%Y} (büro içi kural) · rapor tarihi {o['bugun']:%d.%m.%Y} · {o['beklenen']} belgeden {o['gelen']} tanesi geldi (%{o['tamamlanma']})."])
    ws = wb.create_sheet("Belge Matrisi")
    _baslik(ws, ["Mükellef"] + r["turler"], (32,) + (18,) * len(r["turler"]), "B2")
    renk = {"Geldi": YESIL, "Geç geldi": MAVI, "Bekleniyor": SARI, "Gecikmiş": KIRMIZI}
    for m in sorted(r["mukellefler"], key=lambda m: m["mukellef"]):
        d = [m["belgeler"].get(katla(t), "—") for t in r["turler"]]
        _satir(ws, [m["mukellef"]] + d, dolgu={j: renk.get(x) for j, x in enumerate(d, 2)})
    ws.append([])
    ws.append(["— : bu mükelleften bu belge beklenmiyor."])
    ws = wb.create_sheet("Eksik Belgeler")
    _baslik(ws, ["Mükellef", "Belge Türü", "Durum", "Gecikme (gün)", "Sorumlu"], (32, 26, 13, 11, 18))
    for s in r["satirlar"]:
        if s["durum"] in ("Gecikmiş", "Bekleniyor"):
            _satir(ws, [s["mukellef"], s["belge"], s["durum"], s["gecikme"] or None, s["sorumlu"]], dolgu={3: renk[s["durum"]]})
    ws = wb.create_sheet("Sorumlu Özeti")
    _baslik(ws, ["Sorumlu", "Beklenen Belge", "Eksik Belge", "Eksiği Olan Mükellefler"], (22, 12, 12, 70))
    for s in r["sorumlular"]:
        _satir(ws, [s["sorumlu"], s["beklenen"], s["eksik"] or None, ", ".join(s["mukellefler"])], sar=(4,))
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(beklenen_yol: Path, gelen_yol: Path | None, cikti: Path, donem: tuple[int, int], bugun: date, son_gun: int = 10) -> dict:
    beklenen, gelen, uy = oku(beklenen_yol, gelen_yol)
    r = takip(beklenen, gelen, donem, bugun, son_gun)
    uy += r["uyarilar"]
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Mükelleflerden beklenen belgeleri gelenlerle karşılaştırıp eksik listesini çıkarır.")
    p.add_argument("--beklenen", type=Path, default=ORNEK / "beklenen_belgeler.csv", help="Mükellef, Belge Türü, Sorumlu, İletişim")
    p.add_argument("--gelen", type=Path, help="Mükellef, Belge Türü, Dönem, Geliş Tarihi, Not")
    p.add_argument("--donem", help="Kontrol edilecek dönem (ör. 2026-09 veya 09.2026)")
    p.add_argument("--tarih", help="Bugünün tarihi (GG.AA.YYYY); verilmezse sistem tarihi")
    p.add_argument("--son-gun", type=int, default=10, help="Belgelerin gelmesi gereken, dönemi izleyen ayın günü (varsayılan 10)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "evrak_eksik_takibi.xlsx")
    a = p.parse_args(argv)
    ornek = a.beklenen == ORNEK / "beklenen_belgeler.csv"
    gelen = a.gelen or (ORNEK / "gelen_belgeler.csv" if ornek else None)
    donem = donem_coz(a.donem) if a.donem else (2026, 9) if ornek else None
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 13) if ornek and not a.donem else date.today()
    if donem is None or bugun is None or not 1 <= a.son_gun <= 28:
        print("[X] --donem verilmeli (ör. 2026-09); --tarih GG.AA.YYYY biçiminde, --son-gun 1–28 arasında olmalı")
        return 2
    for y in (a.beklenen, gelen):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.beklenen, gelen, a.cikti, donem, bugun, a.son_gun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {o['donem']} · {o['mukellef']} mükellef · {o['beklenen']} belgeden {o['gelen']} geldi (%{o['tamamlanma']}) · {o['say']['Tamam']} tamam, {o['say']['Eksik var']} eksik var, {o['say']['Hiç gelmedi']} hiç gelmedi")
    for m in r["mukellefler"]:
        if m["eksik"]:
            print(f"[!] {m['mukellef']}: {', '.join(m['eksik'])}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
