import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
VIN = "TESTVN1A2B3C4D5E1"
KOD = {"ariza": {"A100": None}, "iscilik": {"L200": D("1.5")}}


def b(**kw):
    x = {"no": "G1", "sasi": VIN, "bas": date(2025, 3, 15), "ariza": date(2026, 10, 1), "onarim": date(2026, 10, 3), "basvuru": date(2026, 10, 10), "km": D(42000), "akod": "A100", "ikod": "L200", "sure": D("1.5"), "parca": "P1",
         "adet": D(1), "satir": 2}
    x |= kw
    return x | {"ham": {a: "" if x[a] is None else str(x[a]) for a in ("bas", "ariza", "onarim", "basvuru", "km", "sure", "adet")}}


def tek(kodlar=KOD, ayar=None, **kw):
    x = main.kontrol([b(**kw)], kodlar, **(ayar or {}))["basvurular"][0]
    return x["durum"], x["bulgular"]


class KontrolTesti(unittest.TestCase):
    def test_yardimcilar(self):
        self.assertEqual([main.sasi_gecerli(s) for s in (VIN, VIN[:-1], VIN + "1", VIN.replace("E", "O"), VIN.replace("T", "I"), VIN.replace("1", "Q"), VIN.lower(), VIN.replace("A", "-"))], [True] + [False] * 7)
        self.assertEqual((main.ay_ekle(date(2025, 3, 15), 24), main.ay_ekle(date(2024, 2, 29), 12), main.ay_ekle(date(2025, 8, 31), 6), main.ay_ekle(date(2025, 11, 30), 3)),
                         (date(2027, 3, 15), date(2025, 2, 28), date(2026, 2, 28), date(2026, 2, 28)))

    def test_tekil_kurallar(self):
        self.assertEqual(tek(), ("Uygun", []))
        self.assertEqual(tek(sasi="KISA"), ("Ret riski", ["Şasi no hatalı"]))
        self.assertEqual((tek(ariza=date(2027, 3, 15)), tek(ariza=date(2027, 3, 16), onarim=date(2027, 3, 17), basvuru=date(2027, 3, 18)), tek(ariza=date(2025, 3, 14))[1]),
                         (("Düzeltilmeli", ["Tarih sırası hatalı"]), ("Ret riski", ["Garanti süresi dolmuş"]), ["Arıza garanti başlangıcından önce"]))
        self.assertEqual((tek(ayar={"garanti_ay": 12})[1], tek(bas=None)[1], tek(km=None)[1], tek(ayar={"garanti_km": 42000}), tek(ayar={"garanti_km": 41999})[1]),
                         (["Garanti süresi dolmuş"], ["Tarih eksik"], ["Km eksik"], ("Uygun", []), ["Km sınırı aşılmış"]))
        self.assertEqual((tek(basvuru=date(2026, 11, 2)), tek(basvuru=date(2026, 11, 3))[1], tek(basvuru=date(2026, 10, 2))[1], tek(basvuru=date(2026, 10, 6), ayar={"basvuru_gun": 2})[1], tek(basvuru=None), tek(onarim=None)),
                         (("Uygun", []), ["Başvuru süresi geçmiş"], ["Tarih sırası hatalı"], ["Başvuru süresi geçmiş"], ("Uygun", []), ("Uygun", [])))
        self.assertEqual((tek(akod="", ikod=""), tek(akod="X9")[1], tek(ikod="L9")[1], tek(None, akod="X9", ikod="L9", sure=D(9)), tek(sure=D("1.6")), tek(sure=None)[1], tek(sure=D(0))[1]),
                         (("Düzeltilmeli", ["Arıza kodu eksik", "İşçilik kodu eksik"]), ["Tanımsız kod"], ["Tanımsız kod"], ("Uygun", []), ("Düzeltilmeli", ["İşçilik süresi standardı aşıyor"]), ["İşçilik süresi eksik"], ["İşçilik süresi eksik"]))
        self.assertEqual((tek(adet=None)[1], tek(adet=D(0))[1], tek(parca="", adet=None)), (["Parça adedi eksik"], ["Parça adedi eksik"], ("Uygun", [])))

    def test_capraz_kontroller(self):
        l = [b(no="G1"), b(no="G2", ariza=date(2026, 10, 12), onarim=date(2026, 10, 13), basvuru=date(2026, 10, 13), km=D(41000), satir=3), b(no="G3", ariza=date(2026, 11, 1), onarim=None, basvuru=None, km=D(43000), satir=4),
             b(no="G4", ariza=date(2026, 10, 5), onarim=None, basvuru=None, akod="B1", km=D(42100), satir=5), b(no="G1", sasi=VIN[:-1] + "9", satir=6)]
        r = main.kontrol(l, None)
        self.assertEqual([(x["no"], x["durum"], x["bulgular"]) for x in r["basvurular"]],
                         [("G1", "Uygun", []), ("G2", "Ret riski", ["Tekrarlayan onarım", "Km tutarsız"]), ("G3", "Düzeltilmeli", ["Tekrarlayan onarım"]), ("G4", "Uygun", []), ("G1", "Ret riski", ["Mükerrer başvuru no"])])
        self.assertEqual((r["say"], r["turler"]), ({"Ret riski": 2, "Düzeltilmeli": 1, "Uygun": 2}, [("Tekrarlayan onarım", 2), ("Km tutarsız", 1), ("Mükerrer başvuru no", 1)]))
        self.assertEqual(main.kontrol(l, None, tekrar_gun=10)["basvurular"][1]["bulgular"], ["Km tutarsız"])
        with self.assertRaises(ValueError):
            main.kontrol([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "g.xlsx"
            r = main.calistir(ORNEK / "basvurular.csv", cikti, ORNEK / "kodlar.csv")
            self.assertEqual([(x["no"], x["durum"], x["bulgular"]) for x in r["basvurular"]],
                             [("G-001", "Uygun", []), ("G-002", "Ret riski", ["Garanti süresi dolmuş"]), ("G-003", "Ret riski", ["Şasi no hatalı"]), ("G-004", "Ret riski", ["Başvuru süresi geçmiş"]),
                              ("G-005", "Düzeltilmeli", ["İşçilik süresi standardı aşıyor", "Parça adedi eksik"]), ("G-006", "Ret riski", ["Tekrarlayan onarım", "Km tutarsız"]), ("G-007", "Ret riski", ["Tanımsız kod"])])
            self.assertEqual((r["say"], len(r["uyarilar"])), ({"Ret riski": 5, "Düzeltilmeli": 1, "Uygun": 1}, 8))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Kontrol Sonucu", "Bulgu Özeti", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, k = Path(tmp) / "b.csv", Path(tmp) / "k.csv"
            y.write_text("İş Emri No;VIN;Teslim Tarihi;Kabul Tarihi;Kilometre;Arıza Kodu;Operasyon Kodu;Süre\nE1;testvn1a2b3c-4d5e1;15.03.2025;01.10.2026;42.000;a100;l200;1,5\n;TESTVN1A2B3C4D5E2;x;01.10.2026;;;;\n", encoding="utf-8")
            k.write_text("Tür;Kod;Süre\nArıza Kodu;a100;\nİşçilik;L200;1,5\nDiğer;Z1;\n", encoding="utf-8")
            l, kodlar, uy = main.oku(y, k)
            self.assertEqual(([(x["no"], x["sasi"], x["bas"], x["km"], x["akod"], x["sure"]) for x in l], kodlar, [u["tur"] for u in uy]),
                             ([("E1", VIN, date(2025, 3, 15), D(42000), "A100", D("1.5")), ("Satır 3", "TESTVN1A2B3C4D5E2", None, None, "", None)], {"ariza": {"A100": None}, "iscilik": {"L200": D("1.5")}}, ["Kod satırı okunamadı"]))
            r = main.kontrol(l, kodlar)
            self.assertEqual([(x["durum"], x["bulgular"]) for x in r["basvurular"]], [("Uygun", []), ("Ret riski", ["Tarih eksik", "Km eksik", "Arıza kodu eksik", "İşçilik kodu eksik"])])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--basvurular", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--kodlar", str(k), "--garanti-ay", "36", "--garanti-km", "100000", "--basvuru-gun", "15", "--tekrar-gun", "60"]), 0)
            for ek in (["--garanti-ay", "0"], ["--garanti-km", "-1"]):
                self.assertEqual(main.main(arg + ek), 2)
            self.assertEqual(main.main(arg + ["--kodlar", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Başvuru No;Km\nA;1\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
