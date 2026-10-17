"""
Garanti Başvuru Dosyası Kontrolü — Gerisi Bende kod bloğu
Satış Sonrası Hizmetler (Servis) › Garanti Uzmanı

Garanti başvurularını üreticiye göndermeden önce şasi, süre, km, kod ve parça bilgileri yönünden denetler; ret riski taşıyanları işaretler.
  Şasi no            = 17 karakter, yalnız harf ve rakam, I / O / Q harfleri yok
  Garanti süresi     = arıza tarihi, garanti başlangıcı ile başlangıç + --garanti-ay (24) arasında olmalı
  Km sınırı          = --garanti-km verilmişse arıza km'si sınırı aşmamalı (0: sınırsız)
  Başvuru süresi     = başvuru tarihi − onarım tarihi ≤ --basvuru-gun (30)
  Kodlar             = arıza ve işçilik kodu dolu olmalı; kod listesi verilmişse listede bulunmalı; işçilik süresi standart süreyi aşmamalı
  Parça              = parça numarası yazılmışsa adet pozitif olmalı
  Çapraz kontroller  = mükerrer başvuru no · aynı şasi ve arıza kodunun --tekrar-gun (30) içinde tekrarı · aynı şaside km'nin geriye gitmesi
  Durum              = Ret riski (yüksek önemli bulgu) · Düzeltilmeli (orta önemli bulgu) · Uygun
Süre, km ve başvuru kuralları markaya ve sözleşmeye göre değişir; parametreleri üreticinizin garanti el kitabına göre verin. Kod teknik teşhisin doğruluğunu denetlemez.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
import calendar

SUTUN = {"no": ("basvuru no", "basvuru", "talep no", "is emri no"), "sasi": ("sasi no", "sasi", "vin", "sase no"), "bas": ("garanti baslangic", "garanti baslangici", "garanti baslangic tarihi", "tescil tarihi", "teslim tarihi"),
         "ariza": ("ariza tarihi", "arac kabul tarihi", "kabul tarihi"), "onarim": ("onarim tarihi", "onarim bitis", "is emri kapanis"), "basvuru": ("basvuru tarihi", "gonderim tarihi"), "km": ("km", "kilometre"),
         "akod": ("ariza kodu", "hasar kodu", "sikayet kodu"), "ikod": ("iscilik kodu", "operasyon kodu"), "sure": ("iscilik suresi", "iscilik suresi saat", "sure", "sure saat"), "parca": ("parca no", "parca", "parca numarasi"),
         "adet": ("parca adedi", "adet")}
KOD = {"tur": ("kod turu", "tur"), "kod": ("kod",), "sure": ("standart sure", "standart sure saat", "sure")}
DURUM_RENK = {"Ret riski": KIRMIZI, "Düzeltilmeli": SARI, "Uygun": YESIL}


def olcu(x) -> Decimal | None:
    """Süre: 1,5 = 1.5 saat."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    try:
        return Decimal(re.sub(r"[^\d,.\-]", "", str(x)).replace(",", "."))
    except InvalidOperation:
        return None


def ay_ekle(t: date, ay: int) -> date:
    y, a = divmod(t.year * 12 + t.month - 1 + ay, 12)
    return date(y, a + 1, min(t.day, calendar.monthrange(y, a + 1)[1]))


def sasi_gecerli(s: str) -> bool:
    return bool(re.fullmatch(r"[A-HJ-NPR-Z0-9]{17}", s))


def oku(yol: Path, kod_yol: Path | None = None) -> tuple[list[dict], dict | None, list[dict]]:
    uy, l, kodlar = [], [], None
    for n, r in enumerate(kayitlar(yol, SUTUN, ("sasi", "ariza")), 1):
        if not metin(r["sasi"]) and not metin(r["no"]):
            continue
        l.append({"no": metin(r["no"]) or f"Satır {r['_satir']}", "sasi": re.sub(r"[\s\-]", "", metin(r["sasi"])).upper(), "bas": tarih(r["bas"]), "ariza": tarih(r["ariza"]), "onarim": tarih(r["onarim"]), "basvuru": tarih(r["basvuru"]),
                  "km": para(r["km"]) if metin(r["km"]) else None, "akod": metin(r["akod"]).upper(), "ikod": metin(r["ikod"]).upper(), "sure": olcu(r["sure"]) if metin(r["sure"]) else None, "parca": metin(r["parca"]),
                  "adet": para(r["adet"]) if metin(r["adet"]) else None, "ham": {a: metin(r[a]) for a in ("bas", "ariza", "onarim", "basvuru", "km", "sure", "adet")}, "satir": r["_satir"]})
    if kod_yol:
        kodlar = {"ariza": {}, "iscilik": {}}
        for r in kayitlar(kod_yol, KOD, ("tur", "kod")):
            t, k = katla(r["tur"]).split(" ")[0] if metin(r["tur"]) else "", metin(r["kod"]).upper()
            if not k:
                continue
            if t not in kodlar:
                uy.append(uyari("Orta", "Kod satırı okunamadı", f"{k} (satır {r['_satir']})", f"Kod türü '{metin(r['tur'])}' Arıza veya İşçilik olmalı. Satır alınmadı."))
                continue
            kodlar[t][k] = olcu(r["sure"]) if metin(r["sure"]) else None
    return l, kodlar, uy


def kontrol(basvurular: list[dict], kodlar: dict | None = None, garanti_ay: int = 24, garanti_km: int = 0, basvuru_gun: int = 30, tekrar_gun: int = 30) -> dict:
    if not basvurular:
        raise ValueError("Kontrol edilecek başvuru yok.")
    bulgu = {id(b): [] for b in basvurular}

    def ekle(b: dict, onem: str, tur: str, ac: str) -> None:
        bulgu[id(b)].append((onem, tur, ac))

    gorulen, sasi_grup = {}, defaultdict(list)
    for b in basvurular:
        h = b["ham"]
        if b["no"] in gorulen:
            ekle(b, "Yüksek", "Mükerrer başvuru no", f"Aynı başvuru numarası satır {gorulen[b['no']]} ile tekrar ediyor.")
        gorulen.setdefault(b["no"], b["satir"])
        if not sasi_gecerli(b["sasi"]):
            ekle(b, "Yüksek", "Şasi no hatalı", f"'{b['sasi']}' 17 karakter değil veya geçersiz karakter (I, O, Q, işaret) içeriyor.")
        else:
            sasi_grup[b["sasi"]].append(b)
        for alan, ad in (("bas", "Garanti başlangıcı"), ("ariza", "Arıza tarihi")):
            if b[alan] is None:
                ekle(b, "Yüksek", "Tarih eksik", f"{ad} boş veya okunamadı ('{h[alan]}'); garanti süresi kontrol edilemedi.")
        for alan, ad in (("onarim", "Onarım tarihi"), ("basvuru", "Başvuru tarihi")):
            if b[alan] is None and h[alan]:
                ekle(b, "Orta", "Tarih okunamadı", f"{ad} '{h[alan]}' okunamadı.")
        if b["bas"] and b["ariza"]:
            bitis = ay_ekle(b["bas"], garanti_ay)
            if b["ariza"] < b["bas"]:
                ekle(b, "Yüksek", "Arıza garanti başlangıcından önce", f"Arıza {b['ariza']:%d.%m.%Y}, garanti başlangıcı {b['bas']:%d.%m.%Y}. Tarihleri kontrol edin.")
            elif b["ariza"] > bitis:
                ekle(b, "Yüksek", "Garanti süresi dolmuş", f"Garanti {bitis:%d.%m.%Y} tarihinde bitiyor ({garanti_ay} ay); arıza {b['ariza']:%d.%m.%Y}. Uzatılmış garanti veya iyi niyet onayı yoksa reddedilir.")
        if b["km"] is None or b["km"] < 0:
            ekle(b, "Yüksek", "Km eksik", f"Km değeri boş veya okunamadı ('{h['km']}').")
        elif garanti_km and b["km"] > garanti_km:
            ekle(b, "Yüksek", "Km sınırı aşılmış", f"Km {b['km']:,.0f}, garanti sınırı {garanti_km:,}.")
        if b["ariza"] and b["onarim"] and b["onarim"] < b["ariza"]:
            ekle(b, "Orta", "Tarih sırası hatalı", f"Onarım tarihi ({b['onarim']:%d.%m.%Y}) arıza tarihinden ({b['ariza']:%d.%m.%Y}) önce.")
        if b["onarim"] and b["basvuru"]:
            g = (b["basvuru"] - b["onarim"]).days
            if g > basvuru_gun:
                ekle(b, "Yüksek", "Başvuru süresi geçmiş", f"Başvuru onarımdan {g} gün sonra yapılmış (sınır {basvuru_gun} gün).")
            elif g < 0:
                ekle(b, "Orta", "Tarih sırası hatalı", f"Başvuru tarihi ({b['basvuru']:%d.%m.%Y}) onarım tarihinden ({b['onarim']:%d.%m.%Y}) önce.")
        for alan, ad, tur in (("akod", "Arıza kodu", "ariza"), ("ikod", "İşçilik kodu", "iscilik")):
            if not b[alan]:
                ekle(b, "Orta", f"{ad} eksik", f"{ad} boş.")
            elif kodlar is not None and b[alan] not in kodlar[tur]:
                ekle(b, "Yüksek", "Tanımsız kod", f"{ad} '{b[alan]}' kod listesinde yok.")
        if b["ikod"]:
            std = kodlar["iscilik"].get(b["ikod"]) if kodlar else None
            if b["sure"] is None or b["sure"] <= 0:
                ekle(b, "Orta", "İşçilik süresi eksik", f"İşçilik kodu var ama süre boş veya sıfır ('{h['sure']}').")
            elif std is not None and b["sure"] > std:
                ekle(b, "Orta", "İşçilik süresi standardı aşıyor", f"Talep edilen {b['sure']:f} saat, standart {std:f} saat. Ek süre için gerekçe veya onay ekleyin.")
        if b["parca"] and (b["adet"] is None or b["adet"] <= 0):
            ekle(b, "Orta", "Parça adedi eksik", f"Parça {b['parca']} için adet boş veya sıfır ('{h['adet']}').")
    for sasi, grup in sasi_grup.items():
        tarihli = sorted((b for b in grup if b["ariza"]), key=lambda b: (b["ariza"], b["satir"]))
        for i, b in enumerate(tarihli):
            for o in tarihli[:i]:
                if b["akod"] and b["akod"] == o["akod"] and (b["ariza"] - o["ariza"]).days <= tekrar_gun:
                    ekle(b, "Orta", "Tekrarlayan onarım", f"Aynı araçta aynı arıza kodu ({b['akod']}) {(b['ariza'] - o['ariza']).days} gün önce de işlenmiş (başvuru {o['no']}). Üretici ilk onarımı sorgulayabilir.")
                    break
            onceki = [o for o in tarihli[:i] if o["km"] is not None and b["km"] is not None and o["ariza"] < b["ariza"] and o["km"] > b["km"]]
            if onceki:
                ekle(b, "Yüksek", "Km tutarsız", f"Km {b['km']:,.0f}; aynı aracın {onceki[-1]['ariza']:%d.%m.%Y} tarihli başvurusunda {onceki[-1]['km']:,.0f} yazılmış. Km geriye gidemez.")
    liste, uy = [], []
    for b in basvurular:
        e = bulgu[id(b)]
        onem = {o for o, _, _ in e}
        durum = "Ret riski" if "Yüksek" in onem else "Düzeltilmeli" if "Orta" in onem else "Uygun"
        liste.append({k: b[k] for k in ("no", "sasi", "bas", "ariza", "onarim", "basvuru", "km", "akod", "ikod", "sure", "parca", "adet")} | {"durum": durum, "bulgular": [t for _, t, _ in e]})
        uy += [uyari(o, t, b["no"], ac) for o, t, ac in e]
    say = {d: sum(1 for x in liste if x["durum"] == d) for d in DURUM_RENK}
    turler = defaultdict(int)
    for u in uy:
        turler[u["tur"]] += 1
    return {"basvurular": liste, "say": say, "turler": sorted(turler.items(), key=lambda t: (-t[1], t[0])), "ayar": (garanti_ay, garanti_km, basvuru_gun, tekrar_gun), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    ga, gk, bg, tg = r["ayar"]
    ws = wb.active
    ws.title = "Kontrol Sonucu"
    _baslik(ws, ["Durum", "Başvuru No", "Şasi No", "Garanti Başlangıcı", "Arıza Tarihi", "Onarım Tarihi", "Başvuru Tarihi", "Km", "Arıza Kodu", "İşçilik Kodu", "İşçilik Süresi", "Parça No", "Adet", "Bulgular", "Düzeltme Notu (siz doldurun)"],
            (13, 14, 22, 13, 12, 12, 12, 10, 12, 12, 9, 16, 7, 56, 30))
    for b in r["basvurular"]:
        _satir(ws, [b["durum"], b["no"], b["sasi"], b["bas"], b["ariza"], b["onarim"], b["basvuru"], b["km"], b["akod"], b["ikod"], b["sure"], b["parca"], b["adet"], ", ".join(b["bulgular"]), ""], tarih_s=(4, 5, 6, 7),
               dolgu={1: DURUM_RENK[b["durum"]]}, sar=(14,))
        ws.cell(ws.max_row, 8).number_format = "#,##0"
    ws.append([])
    ws.append([f"Kurallar: garanti {ga} ay" + (f", {gk:,} km" if gk else ", km sınırsız") + f"; başvuru onarımdan sonra en geç {bg} gün; aynı arıza {tg} gün içinde tekrarlarsa işaretlenir. Parametreleri üreticinizin garanti kurallarına göre verin."])
    ws = wb.create_sheet("Bulgu Özeti")
    _baslik(ws, ["Bulgu", "Adet"], (40, 10))
    for t, n in r["turler"]:
        _satir(ws, [t, n])
    ws.append([])
    for d, n in r["say"].items():
        _satir(ws, [d, n], dolgu={1: DURUM_RENK[d]})
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, kod_yol: Path | None = None, **ayar) -> dict:
    basvurular, kodlar, uy = oku(yol, kod_yol)
    r = kontrol(basvurular, kodlar, **ayar)
    uy += r["uyarilar"]
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Garanti başvurularını şasi, süre, km, kod ve parça bilgileri yönünden denetler.")
    p.add_argument("--basvurular", type=Path, default=ORNEK / "basvurular.csv", help="Başvuru No, Şasi No, Garanti Başlangıç, Arıza Tarihi, Onarım Tarihi, Başvuru Tarihi, Km, Arıza Kodu, İşçilik Kodu, İşçilik Süresi, Parça No, Parça Adedi")
    p.add_argument("--kodlar", type=Path, help="İsteğe bağlı: Kod Türü (Arıza / İşçilik), Kod, Standart Süre")
    p.add_argument("--garanti-ay", type=int, default=24, help="Garanti süresi, ay (varsayılan 24)")
    p.add_argument("--garanti-km", type=int, default=0, help="Garanti km sınırı; 0 ise sınırsız (varsayılan 0)")
    p.add_argument("--basvuru-gun", type=int, default=30, help="Onarımdan sonra başvuru için azami gün (varsayılan 30)")
    p.add_argument("--tekrar-gun", type=int, default=30, help="Aynı arızanın tekrarı sayılacağı gün aralığı (varsayılan 30)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "garanti_basvuru_kontrol.xlsx")
    a = p.parse_args(argv)
    kod = a.kodlar or (ORNEK / "kodlar.csv" if a.basvurular == ORNEK / "basvurular.csv" else None)
    if a.garanti_ay <= 0 or a.garanti_km < 0 or a.basvuru_gun < 0 or a.tekrar_gun < 0:
        print("[X] --garanti-ay pozitif; --garanti-km, --basvuru-gun ve --tekrar-gun sıfır veya pozitif olmalı")
        return 2
    for y in (a.basvurular, kod):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.basvurular, a.cikti, kod, garanti_ay=a.garanti_ay, garanti_km=a.garanti_km, basvuru_gun=a.basvuru_gun, tekrar_gun=a.tekrar_gun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    s = r["say"]
    print(f"[OK] {len(r['basvurular'])} başvuru · {s['Uygun']} uygun · {s['Düzeltilmeli']} düzeltilmeli · {s['Ret riski']} ret riski")
    for b in r["basvurular"]:
        if b["durum"] == "Ret riski":
            print(f"[!] {b['no']}: {', '.join(b['bulgular'])}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
