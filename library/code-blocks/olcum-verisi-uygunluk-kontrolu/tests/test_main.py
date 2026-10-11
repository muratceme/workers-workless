import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def tol(alt=None, ust=None, nominal=None):
    return {"parca": "P", "kar": "K", "nominal": None if nominal is None else D(nominal), "alt": None if alt is None else D(alt), "ust": None if ust is None else D(ust), "birim": "mm"}


def olc(deger, kar="K", numune="N1", parca="P"):
    return {"parca": parca, "numune": numune, "kar": kar, "deger": D(deger), "tarih": None, "operator": "", "cihaz": ""}


class DegerlendirmeTesti(unittest.TestCase):
    def test_sayi_okuma(self):
        self.assertEqual([main.olcu(x) for x in ("25,000", "25.000", "1.250", "−0,021", "+0,015", 10.5, " 80,02 mm", "", "yok")],
                         [D("25.000"), D("25.000"), D("1.250"), D("-0.021"), D("0.015"), D("10.5"), D("80.02"), None, None])       # nokta binlik ayırıcı değildir

    def test_cift_tarafli(self):
        t = tol("24.979", "25.000", "25.000")
        d = lambda v: main.degerlendir(D(v), t)   # noqa: E731
        self.assertEqual(d("24.990"), ("Uygun", None))
        self.assertEqual(d("25.001"), ("Uygunsuz (üst)", D("0.001")))
        self.assertEqual(d("24.978"), ("Uygunsuz (alt)", D("0.001")))
        self.assertEqual([d(v)[0] for v in ("25.000", "24.979", "24.9811", "24.9812", "24.9979", "24.9978")],
                         ["Sınıra yakın", "Sınıra yakın", "Sınıra yakın", "Uygun", "Sınıra yakın", "Uygun"])                        # aralık 0,021; yüzde 10 = 0,0021
        self.assertEqual(main.degerlendir(D("25.000"), t, D(0))[0], "Sınıra yakın")                                                 # limitin kendisi (uzaklık 0)
        self.assertEqual(main.degerlendir(D("24.999"), t, D(0))[0], "Uygun")

    def test_tek_tarafli(self):
        self.assertEqual([main.degerlendir(D(v), tol(ust="1.6"))[0] for v in ("0.2", "1.59", "1.6", "1.61")], ["Uygun", "Uygun", "Uygun", "Uygunsuz (üst)"])
        self.assertEqual([main.degerlendir(D(v), tol(alt="58"))[0] for v in ("57.9", "58", "65")], ["Uygunsuz (alt)", "Uygun", "Uygun"])

    def test_ozetler(self):
        t = {("p", "k"): tol("9.9", "10.1", "10"), ("p", "ra"): tol(ust="1.6") | {"kar": "Ra"}, ("", "genel"): tol(alt="5") | {"parca": "", "kar": "Genel"}, ("p", "yok"): tol(alt="1") | {"kar": "Yok"}}
        r = main.kontrol([olc("10.0"), olc("10.2", numune="N2"), olc("9.91", numune="N3"), olc("1.7", "Ra", "N3"), olc("7", "Genel", "N1", "Q"), olc("3", "Bilinmeyen")], t)
        self.assertEqual(r["say"], {"Uygun": 2, "Sınıra yakın": 1, "Uygunsuz": 2})
        self.assertEqual({x["numune"]: (x["sonuc"], x["uygunsuz"], x["yakin"]) for x in r["parca"] if x["parca"] == "P"},
                         {"N1": ("KABUL", [], []), "N2": ("RET", ["K"], []), "N3": ("RET", ["Ra"], ["K"])})
        k = next(x for x in r["kar"] if x["kar"] == "K")
        self.assertEqual((k["n"], k["min"], k["max"], k["ort"], k["uygunsuz"], k["oran"], k["yakin"]), (3, D("9.91"), D("10.2"), D("10.0367"), 1, D("33.33"), 1))
        self.assertEqual(sorted(u["tur"] for u in r["uyarilar"]), ["Karakteristik ölçülmemiş", "Ölçümün toleransı tanımsız"])
        self.assertEqual(next(s["sapma"] for s in r["sonuc"] if s["deger"] == D("10.2")), D("0.2"))


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "o.xlsx"
            r = main.calistir(ORNEK / "olcumler.csv", ORNEK / "toleranslar.csv", cikti)
            self.assertEqual((len(r["sonuc"]), r["say"]), (25, {"Uygun": 17, "Sınıra yakın": 3, "Uygunsuz": 5}))
            self.assertEqual({x["numune"]: (x["sonuc"], x["uygunsuz"], x["yakin"]) for x in r["parca"]},
                             {"N-01": ("KABUL", [], []), "N-02": ("RET", ["Boy L1"], ["Çap D1"]), "N-03": ("RET", ["Çap D1", "Delik d2", "Yüzey pürüzlülüğü Ra", "Sertlik"], []),
                              "N-04": ("KABUL", [], []), "N-05": ("KABUL", [], ["Boy L1", "Delik d2"])})
            d1 = next(x for x in r["kar"] if x["kar"] == "Çap D1")
            self.assertEqual((d1["t"]["alt"], d1["t"]["ust"], d1["min"], d1["max"], d1["ort"], d1["uygunsuz"]), (D("24.979"), D("25.000"), D("24.980"), D("25.002"), D("24.9898"), 1))
            self.assertEqual(next(s["tasma"] for s in r["sonuc"] if s["numune"] == "N-02" and s["kar"] == "Boy L1"), D("0.01"))
            self.assertEqual(r["uyarilar"], [])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Parça Özeti", "Karakteristik Özeti", "Ölçüm Sonuçları", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            t, o = Path(tmp) / "t.csv", Path(tmp) / "o.csv"
            t.write_text("Karakteristik;Nominal;Alt Tolerans;Üst Tolerans;Alt Limit;Üst Limit\nA;10;-0,1;0,1;;\nA;10;-0,2;0,2;;\nB;10;0,05;0,05;;\nC;10;;;;\nD;;;;;2,5\nE;10;0,1;0,3;;\n", encoding="utf-8")
            tl, uy = main.tolerans_oku(t)
            self.assertEqual({k[1]: (v["alt"], v["ust"]) for k, v in tl.items()}, {"a": (D("9.9"), D("10.1")), "b": (D("10.05"), D("10.05")), "d": (None, D("2.5")), "e": (D("10.1"), D("10.3"))})
            self.assertEqual(sorted(u["tur"] for u in uy), ["Karakteristik tekrar ediyor", "Tolerans aralığı sıfır", "Tolerans tanımı eksik"])
            o.write_text("Numune No;Karakteristik;Ölçülen\n1;A;10,05\n1;D;yok\n2;a;9,8\n", encoding="utf-8")
            ol, uy = main.olcum_oku(o)
            self.assertEqual(([(x["numune"], x["deger"]) for x in ol], [u["tur"] for u in uy]), ([("1", D("10.05")), ("2", D("9.8"))], ["Ölçüm değeri okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--olcumler", str(o), "--toleranslar", str(t), "--cikti", c]
            self.assertEqual(main.main(arg + ["--sinir", "0"]), 0)
            self.assertEqual(main.main(arg + ["--sinir", "60"]), 2)
            self.assertEqual(main.main(["--olcumler", str(Path(tmp) / "yok.csv")]), 1)
            t.write_text("Karakteristik;Nominal;Alt Tolerans;Üst Tolerans\nA;10;0,2;0,1\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)                                                     # alt limit üstten büyük: kullanılabilir tolerans yok
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
