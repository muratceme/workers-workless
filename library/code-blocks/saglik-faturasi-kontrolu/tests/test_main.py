import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "s.xlsx"
        cls.s = main.calistir(ORNEK / "fatura_kalemleri.csv", ORNEK / "fiyat_listesi.csv", cls.cikti, ORNEK / "provizyonlar.csv",
                              ORNEK / "paket_icerikleri.csv")
        cls.sat = sorted(cls.s["satirlar"], key=lambda x: x.satir)
        cls.oz = {x["fatura"]: x for x in cls.s["ozet"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_satirlar(self):
        self.assertEqual([(x.kod, x.durum, x.kesinti) for x in self.sat[:7]],
                         [("PKT-APP", "Kesinti", D("2000.00")), ("AMH-01", "Kesinti", D("4500.00")), ("YTK-01", "Kesinti", D("5000.00")),
                          ("LAB-HMG", "Kesinti", D("60.00")), ("LAB-CRP", "Uygun", 0), ("LAB-CRP", "Kesinti", D("220.00")), ("ILC-0457", "İnceleme", 0)])
        self.assertIn("Mükerrer", self.sat[5].nedenler[0])
        self.assertIn("paketine dahil", self.sat[1].nedenler[0])
        f502 = self.sat[7:11]
        self.assertEqual([x.durum for x in f502], ["Kesinti", "Uygun", "Kesinti", "İnceleme"])
        self.assertIn("Hesap hatası", f502[2].nedenler[0])
        self.assertIn("provizyon tarihinden", f502[3].nedenler[0])
        ftr = self.sat[11]
        self.assertEqual(ftr.kesinti, D("2300.00"))             # 12 × 50 fiyat farkı + 2 × 850 adet aşımı
        self.assertEqual([x.durum for x in self.sat[12:]], ["Ret", "Beklemede", "Beklemede"])

    def test_ozet(self):
        f = self.oz
        self.assertEqual((f["F-501"]["kabul"], f["F-501"]["sirket"], f["F-501"]["beklemede"]), (D("60520.00"), D("60520.00"), D(380)))
        self.assertEqual((f["F-502"]["kabul"], f["F-502"]["sirket"], f["F-502"]["sigortali"]), (D("1900.00"), D("1520.00"), D("380.00")))
        self.assertEqual((f["F-503"]["sirket"], f["F-503"]["sigortali"]), (D(5500), D("3000.00")))          # 850 katılım + 2.150 provizyon aşımı
        self.assertEqual([f[k]["durum"] for k in ("F-503", "F-504", "F-505", "F-506")], ["Kesintili", "Ret", "Beklemede", "Beklemede"])
        uy = [(u["onem"], u["fatura"]) for u in self.s["uyarilar"]]
        self.assertIn(("Orta", "F-503"), uy)
        self.assertNotIn(("Yüksek", "F-504"), uy)

    def test_excel(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Fatura Özeti", "Satır Kontrolü", "Kesinti Listesi", "Uyarılar"])
        self.assertEqual(wb["Kesinti Listesi"].max_row, 10)
        self.assertEqual(wb["Fatura Özeti"]["O1"].value, "Uzman Onayı")


class KuralTesti(unittest.TestCase):
    def test_genel_fiyat_ve_provizyonsuz(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t / "f.csv").write_text("Fatura No;Kurum;Provizyon No;Hizmet Kodu;Adet;Birim Fiyat (TL);Tutar (TL)\nA;X Hastanesi;P1;K1;2;100;200\n",
                                     encoding="utf-8")
            (t / "l.csv").write_text("Kurum;Hizmet Kodu;Anlaşmalı Fiyat (TL)\n;K1;90\n", encoding="utf-8")
            (t / "p.csv").write_text("Talep No;Karar;Ödenecek\nP1;Onay;1000\n", encoding="utf-8")
            s = main.calistir(t / "f.csv", t / "l.csv", t / "o.xlsx", t / "p.csv")
            self.assertEqual((s["satirlar"][0].kesinti, s["ozet"][0]["sirket"]), (D("20.00"), D("180.00")))

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--fiyatlar", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
