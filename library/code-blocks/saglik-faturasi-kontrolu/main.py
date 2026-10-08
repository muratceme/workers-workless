"""
Sağlık Faturası Kontrolü — Workers / Workless kod bloğu
Sigortacılık › Sağlık Sigortaları › Provizyon Uzmanı

Anlaşmalı kurum faturalarını satır bazında anlaşmalı fiyat listesi, paket içerikleri ve onaylı provizyonla
karşılaştırır; kesinti tutarlarını ve şirket / sigortalı payını hesaplar:
  - Fiyat farkı: birim fiyat > anlaşmalı fiyat → (fark × adet) kesinti
  - Hesap hatası: tutar > adet × birim fiyat → fark kesinti
  - Paket içi hizmet ayrıca faturalanmış (paket koduyla aynı provizyonda) → satırın tamamı kesinti
  - Mükerrer satır (aynı provizyon, kod, tarih, adet ve fiyat) → tekrar eden satır kesinti
  - Adet sınırı (fiyat listesindeki "Maks Adet", provizyon başına) aşıldı → aşan adet kesinti
  - Fiyat listesinde olmayan hizmet; provizyon tarihinden önceki işlem → inceleme
  - Provizyon reddedilmiş → fatura kabul edilmez; provizyon incelemede / bulunamadı → beklemede
  Fatura düzeyinde: şirket payı = kabul edilen tutar × (1 − katılım payı), onaylı provizyon tutarıyla sınırlı.
Rapor: satır kontrolü, fatura özeti, kuruma gönderilecek kesinti listesi, uyarılar. İnternete bağlanmaz.

Kullanım:
    python main.py                                                  # örnek: 6 fatura, 15 satır
    python main.py --faturalar fatura_kalemleri.xlsx --fiyatlar fiyat_listesi.xlsx --provizyonlar provizyonlar.xlsx --paketler paket_icerikleri.csv
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
K2 = Decimal("0.01")

FATURA_SUTUNLARI = {"fatura": ("fatura no", "fatura"), "kurum": ("kurum", "saglik kurumu"), "provizyon": ("provizyon no", "provizyon", "talep no"),
                    "tarih": ("islem tarihi", "hizmet tarihi", "tarih"), "kod": ("hizmet kodu", "kod", "islem kodu"), "ad": ("hizmet adi", "hizmet", "aciklama"),
                    "adet": ("adet", "miktar"), "birim": ("birim fiyat tl", "birim fiyat"), "tutar": ("tutar tl", "tutar", "toplam")}
FIYAT_SUTUNLARI = {"kurum": ("kurum", "kurum grubu"), "kod": ("hizmet kodu", "kod"), "ad": ("hizmet adi", "hizmet"),
                   "fiyat": ("anlasmali fiyat tl", "anlasmali fiyat", "fiyat"), "maks": ("maks adet", "azami adet", "adet siniri")}
PROVIZYON_SUTUNLARI = {"no": ("talep no", "provizyon no", "no"), "tarih": ("talep tarihi", "provizyon tarihi", "tarih"),
                       "sigortali": ("sigortali no", "sigortali"), "karar": ("on karar", "karar", "uzman karari", "durum"),
                       "onay": ("odenecek", "onayli tutar", "onay tutari", "sirket payi"), "katilim": ("katilim payi", "katilim payi yuzde", "katilim")}
PAKET_SUTUNLARI = {"paket": ("paket kodu", "paket"), "dahil": ("dahil hizmet kodu", "dahil kod", "hizmet kodu")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("₺", "").replace(" ", "").replace("%", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tl(x: Decimal | None) -> str:
    return "—" if x is None else f"{x.quantize(K2, ROUND_HALF_UP):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Satir:
    satir: int
    fatura: str
    kurum: str
    provizyon: str
    tarih: date | None
    kod: str
    ad: str
    adet: Decimal
    birim: Decimal
    tutar: Decimal
    anlasmali: Decimal | None = None
    kesinti: Decimal = Decimal(0)
    durum: str = "Uygun"
    nedenler: list = field(default_factory=list)

    @property
    def kabul(self) -> Decimal:
        return Decimal(0) if self.durum in ("Beklemede", "Ret") else self.tutar - self.kesinti


def oku_fatura(yol: Path) -> list[Satir]:
    sonuc = []
    for r in kayitlar(yol, FATURA_SUTUNLARI, ("fatura", "kod")):
        if not r.get("fatura"):
            continue
        adet = para(r.get("adet")) or Decimal(1)
        birim = para(r.get("birim"))
        tutar = para(r.get("tutar"))
        birim = birim if birim is not None else (tutar / adet if tutar is not None else Decimal(0))
        sonuc.append(Satir(r["_satir"], metin(r["fatura"]), metin(r.get("kurum")), metin(r.get("provizyon")), tarih(r.get("tarih")), metin(r["kod"]).upper(),
                           metin(r.get("ad")), adet, birim, tutar if tutar is not None else adet * birim))
    return sonuc


def oku_fiyatlar(yol: Path) -> dict[tuple[str, str], dict]:
    sonuc = {}
    for r in kayitlar(yol, FIYAT_SUTUNLARI, ("kod", "fiyat")):
        if r.get("kod") and para(r.get("fiyat")) is not None:
            maks = para(r.get("maks"))
            sonuc[(katla(r.get("kurum")), metin(r["kod"]).upper())] = {"ad": metin(r.get("ad")), "fiyat": para(r["fiyat"]), "maks": maks}
    return sonuc


def oku_provizyonlar(yol: Path | None) -> dict[str, dict]:
    if yol is None:
        return {}
    sonuc = {}
    for r in kayitlar(yol, PROVIZYON_SUTUNLARI, ("no",)):
        if r.get("no"):
            k = para(r.get("katilim")) or Decimal(0)
            sonuc[metin(r["no"])] = {"tarih": tarih(r.get("tarih")), "sigortali": metin(r.get("sigortali")), "karar": metin(r.get("karar")) or "Onay",
                                     "onay": para(r.get("onay")), "katilim": k / 100 if k > 1 else k}
    return sonuc


def oku_paketler(yol: Path | None) -> dict[str, set[str]]:
    sonuc = defaultdict(set)
    if yol is None:
        return sonuc
    for r in kayitlar(yol, PAKET_SUTUNLARI, ("paket", "dahil")):
        if r.get("paket") and r.get("dahil"):
            sonuc[metin(r["paket"]).upper()].add(metin(r["dahil"]).upper())
    return sonuc


# ----------------------------------------------------------------------------
# Kontrol
# ----------------------------------------------------------------------------

SIRA = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}


def kontrol_et(satirlar: list[Satir], fiyatlar: dict, provizyonlar: dict, paketler: dict) -> dict:
    uyarilar = []

    def fiyat_bul(s: Satir):
        return fiyatlar.get((katla(s.kurum), s.kod)) or fiyatlar.get(("", s.kod))

    def kes(s: Satir, tutar: Decimal, neden: str):
        tutar = min(tutar, s.tutar - s.kesinti)
        if tutar > 0:
            s.kesinti += tutar.quantize(K2, ROUND_HALF_UP)
            s.nedenler.append(f"{neden}: {tl(tutar)} TL")
            s.durum = "Kesinti"

    gorulen = Counter()
    adet_toplam = defaultdict(Decimal)
    prov_kodlar = defaultdict(set)
    for s in satirlar:
        prov_kodlar[s.provizyon].add(s.kod)
    for s in sorted(satirlar, key=lambda s: (s.provizyon, s.tarih or date.min, s.satir)):
        p = provizyonlar.get(s.provizyon)
        karar = katla(p["karar"]) if p else ""
        if p is None:
            s.durum = "Beklemede"
            s.nedenler.append(f"Provizyon bulunamadı ({s.provizyon or 'boş'})")
            continue
        if karar.startswith("ret") or karar.startswith("red"):
            s.durum = "Ret"
            s.kesinti = s.tutar
            s.nedenler.append("Provizyon reddedilmiş; satır kabul edilmez")
            continue
        if karar.startswith("incele") or karar.startswith("bekle"):
            s.durum = "Beklemede"
            s.nedenler.append("Provizyon incelemede; uzman kararı bekleniyor")
            continue
        # Hesap hatası
        hesap = (s.adet * s.birim).quantize(K2, ROUND_HALF_UP)
        if s.tutar > hesap:
            kes(s, s.tutar - hesap, f"Hesap hatası (adet × birim = {tl(hesap)} TL)")
        # Mükerrer
        anahtar = (s.provizyon, s.kod, s.tarih, s.adet, s.birim)
        gorulen[anahtar] += 1
        if gorulen[anahtar] > 1:
            kes(s, s.tutar, "Mükerrer satır")
            continue
        # Paket içi
        paket = next((pk for pk, dahil in paketler.items() if s.kod in dahil and pk in prov_kodlar[s.provizyon]), None)
        if paket:
            kes(s, s.tutar, f"{paket} paketine dahil hizmet ayrıca faturalanmış")
            continue
        f = fiyat_bul(s)
        if f is None:
            s.durum = "İnceleme"
            s.nedenler.append("Anlaşmalı fiyat listesinde yok (ilaç / malzeme ise fatura ve barkod kontrolü)")
            continue
        s.anlasmali = f["fiyat"]
        if s.tarih and p["tarih"] and s.tarih < p["tarih"]:
            s.durum = "İnceleme"
            s.nedenler.append(f"İşlem tarihi ({s.tarih:%d.%m.%Y}) provizyon tarihinden ({p['tarih']:%d.%m.%Y}) önce; bu provizyon kapsamında mı?")
            continue
        if s.birim > f["fiyat"]:
            kes(s, (s.birim - f["fiyat"]) * s.adet, f"Fiyat farkı (birim {tl(s.birim)} > anlaşmalı {tl(f['fiyat'])})")
        if f["maks"]:
            adet_toplam[(s.provizyon, s.kod)] += s.adet
            asan = min(s.adet, adet_toplam[(s.provizyon, s.kod)] - f["maks"])
            if asan > 0:
                kes(s, asan * min(s.birim, f["fiyat"]), f"Adet sınırı aşıldı (provizyon başına en çok {f['maks']:g}, toplam "
                                                        f"{adet_toplam[(s.provizyon, s.kod)]:g})")

    # Fatura özeti
    ozet = []
    gruplar = defaultdict(list)
    for s in satirlar:
        gruplar[(s.fatura, s.provizyon)].append(s)
    for (fno, prov), lst in sorted(gruplar.items()):
        p = provizyonlar.get(prov)
        toplam = sum((s.tutar for s in lst), Decimal(0))
        kesinti = sum((s.kesinti for s in lst if s.durum != "Beklemede"), Decimal(0))
        beklemede = sum((s.tutar for s in lst if s.durum in ("Beklemede", "İnceleme")), Decimal(0))
        kabul = sum((s.kabul for s in lst if s.durum != "İnceleme"), Decimal(0))
        katilim = p["katilim"] if p else Decimal(0)
        sirket = (kabul * (1 - katilim)).quantize(K2, ROUND_HALF_UP)
        onay = p["onay"] if p else None
        limit_asim = Decimal(0)
        if onay is not None and sirket > onay and not katla(p["karar"]).startswith(("ret", "red")):
            limit_asim = sirket - onay
            sirket = onay
            uyarilar.append({"onem": "Orta", "fatura": fno, "aciklama": f"Şirket payı onaylı provizyon tutarını {tl(limit_asim)} TL aşıyor; aşan kısım "
                                                                        "sigortalı payına eklendi (ek provizyon talebi gerekebilir)"})
        durum = ("Ret" if p and katla(p["karar"]).startswith(("ret", "red")) else "Beklemede" if beklemede == toplam
                 else "Kısmi beklemede" if beklemede else "Kesintili" if kesinti else "Uygun")
        ozet.append({"fatura": fno, "kurum": lst[0].kurum, "provizyon": prov, "karar": p["karar"] if p else "Bulunamadı", "satir": len(lst),
                     "toplam": toplam, "kesinti": kesinti, "beklemede": beklemede, "kabul": kabul, "katilim": katilim, "sirket": sirket,
                     "sigortali": kabul - sirket, "onay": onay, "durum": durum})
        if durum != "Ret" and kesinti and toplam and kesinti / toplam >= Decimal("0.1"):
            uyarilar.append({"onem": "Yüksek", "fatura": fno, "aciklama": f"Kesinti oranı %{kesinti / toplam * 100:.1f} ({tl(kesinti)} / {tl(toplam)} TL); "
                                                                          "kurumla mutabakat yapın".replace(".", ",", 1)})
    uyarilar.sort(key=lambda u: (SIRA[u["onem"]], u["fatura"]))
    return {"satirlar": satirlar, "ozet": ozet, "uyarilar": uyarilar}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Uygun": "E3F4E1", "Kesinti": "FDE2E1", "Ret": "FDE2E1", "İnceleme": "FFF4CE", "Beklemede": "FFF4CE", "Kesintili": "FFF4CE",
        "Kısmi beklemede": "FFF4CE", "Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KONTROL = PatternFill("solid", fgColor="FFF4CE")
UST = Alignment(vertical="top", wrap_text=True)
PF = "#,##0.00"


def _baslik(ws, basliklar, genislik):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def rapor_yaz(cikti: Path, s: dict) -> None:
    wb = Workbook()
    oz = wb.active
    oz.title = "Fatura Özeti"
    _baslik(oz, ["Fatura", "Kurum", "Provizyon", "Provizyon Kararı", "Satır", "Fatura Tutarı", "Kesinti", "Beklemede / İnceleme", "Kabul Edilen",
                 "Katılım %", "Şirket Payı", "Sigortalı Payı", "Onaylı Provizyon", "Durum", "Uzman Onayı"],
            (8, 18, 9, 12, 6, 13, 12, 13, 13, 8, 13, 12, 13, 14, 14))
    for x in s["ozet"]:
        oz.append([x["fatura"], x["kurum"], x["provizyon"], x["karar"], x["satir"], float(x["toplam"]), float(x["kesinti"]), float(x["beklemede"]),
                   float(x["kabul"]), float(x["katilim"]), float(x["sirket"]), float(x["sigortali"]), x["onay"] and float(x["onay"]), x["durum"], ""])
        r = oz.max_row
        for j in (6, 7, 8, 9, 11, 12, 13):
            oz.cell(r, j).number_format = PF
        oz.cell(r, 10).number_format = "0%"
        oz.cell(r, 14).fill = PatternFill("solid", fgColor=RENK.get(x["durum"], "FFFFFF"))
        oz.cell(r, 15).fill = KONTROL
    oz.append([])
    t = lambda k: float(sum((x[k] for x in s["ozet"]), Decimal(0)))  # noqa: E731
    oz.append(["Toplam", "", "", "", sum(x["satir"] for x in s["ozet"]), t("toplam"), t("kesinti"), t("beklemede"), t("kabul"), None, t("sirket"),
               t("sigortali")])
    for j in (6, 7, 8, 9, 11, 12):
        oz.cell(oz.max_row, j).number_format = PF
        oz.cell(oz.max_row, j).font = Font(bold=True)

    sk = wb.create_sheet("Satır Kontrolü")
    _baslik(sk, ["Fatura", "Provizyon", "Tarih", "Kod", "Hizmet", "Adet", "Birim Fiyat", "Anlaşmalı Fiyat", "Tutar", "Kesinti", "Kabul", "Durum", "Neden"],
            (8, 9, 11, 10, 30, 6, 11, 11, 11, 11, 11, 11, 60))
    for x in sorted(s["satirlar"], key=lambda x: (x.fatura, x.satir)):
        sk.append([x.fatura, x.provizyon, x.tarih, x.kod, x.ad, float(x.adet), float(x.birim), x.anlasmali and float(x.anlasmali), float(x.tutar),
                   float(x.kesinti), float(x.kabul), x.durum, "\n".join(x.nedenler)])
        r = sk.max_row
        sk.cell(r, 3).number_format = "DD.MM.YYYY"
        for j in (7, 8, 9, 10, 11):
            sk.cell(r, j).number_format = PF
        sk.cell(r, 12).fill = PatternFill("solid", fgColor=RENK[x.durum])
        sk.cell(r, 13).alignment = UST
    sk.auto_filter.ref = f"A1:M{sk.max_row}"

    kl = wb.create_sheet("Kesinti Listesi")
    _baslik(kl, ["Fatura", "Provizyon", "Tarih", "Kod", "Hizmet", "Faturalanan", "Kesinti", "Kesinti Gerekçesi", "Kurum İtirazı"],
            (8, 9, 11, 10, 30, 12, 12, 60, 24))
    for x in sorted(s["satirlar"], key=lambda x: (x.fatura, x.satir)):
        if x.kesinti and x.durum in ("Kesinti", "Ret", "İnceleme"):
            kl.append([x.fatura, x.provizyon, x.tarih, x.kod, x.ad, float(x.tutar), float(x.kesinti), "\n".join(x.nedenler), ""])
            kl.cell(kl.max_row, 3).number_format = "DD.MM.YYYY"
            kl.cell(kl.max_row, 6).number_format = kl.cell(kl.max_row, 7).number_format = PF
            kl.cell(kl.max_row, 8).alignment = UST
            kl.cell(kl.max_row, 9).fill = KONTROL

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Fatura", "Açıklama", "İnceleme"], (9, 8, 90, 24))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["fatura"], u["aciklama"], ""])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).fill = KONTROL
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(fatura_yolu: Path, fiyat_yolu: Path, cikti: Path, provizyon_yolu: Path | None = None, paket_yolu: Path | None = None) -> dict:
    s = kontrol_et(oku_fatura(fatura_yolu), oku_fiyatlar(fiyat_yolu), oku_provizyonlar(provizyon_yolu), oku_paketler(paket_yolu))
    rapor_yaz(cikti, s)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Kurum faturalarını anlaşmalı fiyat listesi ve onaylı provizyonla satır bazında karşılaştırır.")
    p.add_argument("--faturalar", type=Path, default=ORNEK / "fatura_kalemleri.csv")
    p.add_argument("--fiyatlar", type=Path, default=ORNEK / "fiyat_listesi.csv", help="Kurum, Hizmet Kodu, Anlaşmalı Fiyat, Maks Adet")
    p.add_argument("--provizyonlar", type=Path, help="Talep No, Talep Tarihi, Karar, Ödenecek, Katılım Payı %% (Provizyon Talebi Değerlendirme çıktısı)")
    p.add_argument("--paketler", type=Path, help="Paket Kodu, Dahil Hizmet Kodu")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "saglik_fatura_kontrolu.xlsx")
    a = p.parse_args(argv)
    if a.faturalar == ORNEK / "fatura_kalemleri.csv":
        a.provizyonlar = a.provizyonlar or ORNEK / "provizyonlar.csv"
        a.paketler = a.paketler or ORNEK / "paket_icerikleri.csv"
    for y in (a.faturalar, a.fiyatlar, a.provizyonlar, a.paketler):
        if y is not None and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        s = calistir(a.faturalar, a.fiyatlar, a.cikti, a.provizyonlar, a.paketler)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    for x in s["ozet"]:
        print(f"[{'OK' if x['durum'] == 'Uygun' else '!'}] {x['fatura']} ({x['provizyon']}): fatura {tl(x['toplam'])} · kesinti {tl(x['kesinti'])} · "
              f"şirket payı {tl(x['sirket'])} TL · {x['durum']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
