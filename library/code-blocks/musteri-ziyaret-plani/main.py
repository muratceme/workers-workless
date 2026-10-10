"""
Müşteri Ziyaret Planı — Workers / Workless kod bloğu
Satış › Satış Temsilcisi

Müşteri listesi, konum ve ziyaret sıklığına göre haftalık ziyaret planı ve rota sırası çıkarır:
  - Bu hafta ziyareti gelenler: son ziyaret + sıklık (haftalık 7, iki haftada 14, aylık 28, üç ayda 84 gün) hafta sonuna
    kadar doluyorsa; hiç ziyaret edilmemiş müşteri de gelir.
  - Günlere dağıtım: müşteriler başlangıç noktasına göre açı sırasına dizilir (süpürme yöntemi) ve günlük süre
    kapasitesine göre ardışık gruplara bölünür; böylece aynı gün birbirine yakın müşteriler ziyaret edilir. "Uygun
    Günler" kısıtı olan müşteri, izinli günlerden rotayı en az uzatanına eklenir.
  - Gün içi sıra: en yakın komşu + 2-opt iyileştirme (başlangıç noktasından çıkış ve dönüş).
  - Mesafe: kuş uçuşu (haversine) × yol katsayısı (1,3); süre = mesafe / ortalama hız + ziyaret süresi.
  - Günlük süre aşılırsa önceliği en düşük, gecikmesi en az olan müşteri çıkarılır ve sonraki haftaya bırakılır.
Rapor: haftalık plan (gün, sıra, tahmini varış), gün özeti, sığmayanlar, bu hafta gerekmeyenler, uyarılar.
İnternete bağlanmaz; harita servisi kullanmaz.

Kullanım:
    python main.py                                                  # örnek: 48 müşteri, hafta 12.10.2026
    python main.py --musteriler musteriler.xlsx --hafta 12.10.2026 --baslangic 40.99,29.12 --gunluk-dk 480 --hiz 30
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
GUNLER = ["Pazartesi", "Salı", "Çarşamba", "Perşembe", "Cuma", "Cumartesi"]
GUN_KISA = {"pzt": 0, "pazartesi": 0, "sal": 1, "sali": 1, "car": 2, "carsamba": 2, "per": 3, "persembe": 3, "cum": 4, "cuma": 4, "cmt": 5, "cumartesi": 5}
ONCELIK = {"a": 0, "b": 1, "c": 2}

SUTUNLAR = {"no": ("musteri no", "musteri kodu", "cari kod"), "ad": ("musteri", "musteri adi", "unvan"), "ilce": ("ilce", "sehir", "bolge", "adres"),
            "enlem": ("enlem", "lat", "latitude"), "boylam": ("boylam", "lon", "lng", "longitude"), "siklik": ("ziyaret sikligi", "siklik", "periyot"),
            "sure": ("ziyaret suresi", "ziyaret suresi dk", "sure"), "oncelik": ("oncelik", "sinif", "segment"), "son": ("son ziyaret", "son ziyaret tarihi"),
            "gunler": ("uygun gunler", "ziyaret gunleri")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def ondalik(x) -> float | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return float(x)
    try:
        return float(str(x).strip().replace(",", "."))
    except ValueError:
        return None


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def siklik_gun(x) -> int | None:
    """'Haftalık' → 7, '2 haftada bir' → 14, 'Aylık' → 28, '3 ayda bir' → 84, '10' → 10 (gün)."""
    k = katla(x)
    if not k:
        return None
    m = re.search(r"(\d+)", k)
    n = int(m[1]) if m else 1
    if "hafta" in k:
        return 7 * n
    if "ay" in k.split() or "aylik" in k or "ayda" in k:
        return 28 * n
    if "gun" in k or k.isdigit():
        return n
    return None


def gunleri_coz(x) -> set[int]:
    return {GUN_KISA[p] for p in re.split(r"[ ,;/]+", katla(x).replace(" ", ",")) if p in GUN_KISA}


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


# ----------------------------------------------------------------------------
# Geometri ve rota
# ----------------------------------------------------------------------------

def haversine(a: tuple[float, float], b: tuple[float, float]) -> float:
    """İki nokta arası kuş uçuşu mesafe (km)."""
    la1, lo1, la2, lo2 = map(math.radians, (a[0], a[1], b[0], b[1]))
    h = math.sin((la2 - la1) / 2) ** 2 + math.cos(la1) * math.cos(la2) * math.sin((lo2 - lo1) / 2) ** 2
    return 2 * 6371.0088 * math.asin(math.sqrt(h))


def rota_km(baslangic, noktalar, katsayi: float) -> float:
    if not noktalar:
        return 0.0
    yol = [baslangic] + noktalar + [baslangic]
    return sum(haversine(yol[i], yol[i + 1]) for i in range(len(yol) - 1)) * katsayi


def rota_sirala(baslangic, musteriler: list) -> list:
    """En yakın komşu + 2-opt (kapalı tur: başlangıçtan çıkıp başlangıca dönüş)."""
    kalan, sira, konum = list(musteriler), [], baslangic
    while kalan:
        en = min(kalan, key=lambda m: (haversine(konum, m.konum), m.no))
        sira.append(en)
        kalan.remove(en)
        konum = en.konum
    def uzunluk(s):
        return rota_km(baslangic, [m.konum for m in s], 1.0)
    iyilesti = True
    while iyilesti and len(sira) >= 3:
        iyilesti = False
        for i in range(len(sira) - 1):
            for j in range(i + 1, len(sira)):
                yeni = sira[:i] + sira[i:j + 1][::-1] + sira[j + 1:]
                if uzunluk(yeni) < uzunluk(sira) - 1e-9:
                    sira, iyilesti = yeni, True
    return sira


# ----------------------------------------------------------------------------
# Veri ve plan
# ----------------------------------------------------------------------------

@dataclass
class Musteri:
    no: str
    ad: str
    ilce: str
    konum: tuple[float, float] | None
    siklik: int
    sure: float
    oncelik: str
    son: date | None
    gunler: set[int] = field(default_factory=set)
    vade: date | None = None
    gecikme: int = 0


def oku(yol: Path, varsayilan_sure: float) -> tuple[list[Musteri], list[dict]]:
    lst, uy = [], []
    for r in kayitlar(yol, SUTUNLAR, ("no",)):
        no = metin(r.get("no"))
        if not no:
            continue
        la, lo = ondalik(r.get("enlem")), ondalik(r.get("boylam"))
        konum = (la, lo) if la is not None and lo is not None and -90 <= la <= 90 and -180 <= lo <= 180 else None
        sk = siklik_gun(r.get("siklik"))
        if sk is None:
            sk = 28
            uy.append({"onem": "Bilgi", "tur": "Sıklık okunamadı", "kim": no, "aciklama": f"'{metin(r.get('siklik'))}'; aylık (28 gün) kabul edildi"})
        lst.append(Musteri(no, metin(r.get("ad")), metin(r.get("ilce")), konum, sk, ondalik(r.get("sure")) or varsayilan_sure,
                           (metin(r.get("oncelik")) or "B").upper()[:1], tarih(r.get("son")), gunleri_coz(r.get("gunler"))))
    return lst, uy


def gun_suresi(baslangic, musteriler: list[Musteri], hiz: float, katsayi: float) -> tuple[float, float]:
    km = rota_km(baslangic, [m.konum for m in musteriler], katsayi)
    return km, km / hiz * 60 + sum(m.sure for m in musteriler)


def planla(musteriler: list[Musteri], hafta: date, baslangic: tuple[float, float], *, gun_sayisi: int = 5, gunluk_dk: float = 480, hiz: float = 30,
           katsayi: float = 1.3, baslama: str = "09:00") -> dict:
    uy = []
    hafta_sonu = hafta + timedelta(days=gun_sayisi - 1)
    gelen, gerekmeyen, konumsuz = [], [], []
    for m in musteriler:
        m.vade = m.son + timedelta(days=m.siklik) if m.son else hafta
        m.gecikme = (hafta - m.vade).days
        if m.vade > hafta_sonu:
            gerekmeyen.append(m)
        elif m.konum is None:
            konumsuz.append(m)
            uy.append({"onem": "Orta", "tur": "Konum yok", "kim": f"{m.no} {m.ad}", "aciklama": "Enlem / boylam yok veya geçersiz; plana alınamadı"})
        else:
            gelen.append(m)
            if m.gecikme > m.siklik:
                uy.append({"onem": "Orta", "tur": "Ziyaret gecikmiş", "kim": f"{m.no} {m.ad}",
                           "aciklama": f"Son ziyaret {m.son:%d.%m.%Y}; {m.gecikme} gün gecikme (sıklık {m.siklik} gün)"})

    def aci(m):
        return math.atan2(m.konum[0] - baslangic[0], (m.konum[1] - baslangic[1]) * math.cos(math.radians(baslangic[0])))

    serbest = sorted((m for m in gelen if not m.gunler or len(m.gunler & set(range(gun_sayisi))) == gun_sayisi), key=lambda m: (aci(m), m.no))
    kisitli = [m for m in gelen if m not in serbest]
    gunler: list[list[Musteri]] = [[] for _ in range(gun_sayisi)]
    # Süpürme: hedef gün yükü = toplam ziyaret süresi / gün (yol payı için kapasitenin %75'i üst sınır)
    toplam = sum(m.sure for m in gelen)
    hedef = min(max(toplam / gun_sayisi, 1), gunluk_dk * 0.75)
    g, yuk = 0, 0.0
    for m in serbest:
        if yuk + m.sure > hedef and g < gun_sayisi - 1 and gunler[g]:
            g, yuk = g + 1, 0.0
        gunler[g].append(m)
        yuk += m.sure
    sigmayan = []
    for m in sorted(kisitli, key=lambda m: (ONCELIK.get(m.oncelik.lower(), 1), -m.gecikme, m.no)):
        izinli = sorted(d for d in m.gunler if d < gun_sayisi)
        if not izinli:
            sigmayan.append((m, "Uygun günleri plan günlerinin dışında"))
            continue
        en = min(izinli, key=lambda d: (rota_km(baslangic, [x.konum for x in rota_sirala(baslangic, gunler[d] + [m])], katsayi)
                                         - rota_km(baslangic, [x.konum for x in rota_sirala(baslangic, gunler[d])], katsayi), d))
        gunler[en].append(m)
    # Rota ve kapasite
    plan = []
    for d in range(gun_sayisi):
        sira = rota_sirala(baslangic, gunler[d])
        while sira:
            km, dk = gun_suresi(baslangic, sira, hiz, katsayi)
            if dk <= gunluk_dk:
                break
            cikan = max(sira, key=lambda m: (ONCELIK.get(m.oncelik.lower(), 1), -m.gecikme, m.no))
            sira.remove(cikan)
            sigmayan.append((cikan, f"{GUNLER[d]} günlük süre ({gunluk_dk:g} dk) aşıldı"))
            sira = rota_sirala(baslangic, sira)
        km, dk = gun_suresi(baslangic, sira, hiz, katsayi) if sira else (0.0, 0.0)
        saat = datetime.combine(hafta + timedelta(days=d), datetime.strptime(baslama, "%H:%M").time())
        onceki, satirlar = baslangic, []
        for i, m in enumerate(sira, 1):
            mesafe = haversine(onceki, m.konum) * katsayi
            saat += timedelta(minutes=mesafe / hiz * 60)
            satirlar.append({"sira": i, "m": m, "km": mesafe, "varis": saat, "ayrilis": saat + timedelta(minutes=m.sure)})
            saat += timedelta(minutes=m.sure)
            onceki = m.konum
        donus = haversine(onceki, baslangic) * katsayi if sira else 0.0
        plan.append({"gun": d, "tarih": hafta + timedelta(days=d), "satirlar": satirlar, "km": km, "dk": dk, "donus_km": donus,
                     "bitis": saat + timedelta(minutes=donus / hiz * 60) if sira else None})
    for m, neden in sigmayan:
        uy.append({"onem": "Yüksek" if m.oncelik == "A" else "Orta", "tur": "Plana sığmadı", "kim": f"{m.no} {m.ad}", "aciklama": f"{neden}; sonraki haftaya bırakıldı"})
    sira_ = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    uy.sort(key=lambda u: (sira_[u["onem"]], u["tur"], u["kim"]))
    return {"plan": plan, "sigmayan": sigmayan, "gerekmeyen": sorted(gerekmeyen, key=lambda m: m.vade), "konumsuz": konumsuz, "uyarilar": uy, "hafta": hafta,
            "baslangic": baslangic, "parametre": {"gunluk_dk": gunluk_dk, "hiz": hiz, "katsayi": katsayi}, "gelen": len(gelen) + len(konumsuz)}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
UST = Alignment(vertical="top", wrap_text=True)


def _baslik(ws, basliklar, genislik=()):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def harita_baglantisi(konum) -> str:
    return f"https://www.google.com/maps/search/?api=1&query={konum[0]:.6f},{konum[1]:.6f}"


def rapor_yaz(cikti: Path, s: dict) -> None:
    wb = Workbook()
    hp = wb.active
    hp.title = "Haftalık Plan"
    _baslik(hp, ["Gün", "Tarih", "Sıra", "Müşteri No", "Müşteri", "İlçe", "Öncelik", "Tahmini Varış", "Ziyaret (dk)", "Önceki Noktadan (km)", "Son Ziyaret",
                 "Gecikme (gün)", "Harita", "Ziyaret Notu"], (11, 11, 5, 11, 28, 14, 8, 9, 8, 10, 11, 9, 14, 30))
    for g in s["plan"]:
        for x in g["satirlar"]:
            m = x["m"]
            hp.append([GUNLER[g["gun"]], g["tarih"], x["sira"], m.no, m.ad, m.ilce, m.oncelik, x["varis"].strftime("%H:%M"), m.sure, round(x["km"], 1), m.son,
                       max(m.gecikme, 0), "Haritada aç", ""])
            hp.cell(hp.max_row, 2).number_format = hp.cell(hp.max_row, 11).number_format = "DD.MM.YYYY"
            hp.cell(hp.max_row, 13).hyperlink = harita_baglantisi(m.konum)
            hp.cell(hp.max_row, 13).font = Font(color="0563C1", underline="single")
            hp.cell(hp.max_row, 14).fill = PatternFill("solid", fgColor="FFF4CE")
            if m.gecikme > m.siklik:
                hp.cell(hp.max_row, 12).fill = PatternFill("solid", fgColor="FDE2E1")
    hp.auto_filter.ref = f"A1:N{hp.max_row}"

    go = wb.create_sheet("Gün Özeti")
    _baslik(go, ["Gün", "Tarih", "Ziyaret", "Toplam Yol (km)", "Ziyaret Süresi (dk)", "Yol Süresi (dk)", "Toplam (dk)", "Kapasite Kullanımı", "Tahmini Bitiş"],
            (11, 11, 8, 10, 10, 10, 10, 10, 9))
    p = s["parametre"]
    for g in s["plan"]:
        ziyaret = sum(x["m"].sure for x in g["satirlar"])
        go.append([GUNLER[g["gun"]], g["tarih"], len(g["satirlar"]), round(g["km"], 1), round(ziyaret), round(g["dk"] - ziyaret), round(g["dk"]),
                   g["dk"] / p["gunluk_dk"], g["bitis"].strftime("%H:%M") if g["bitis"] else ""])
        go.cell(go.max_row, 2).number_format = "DD.MM.YYYY"
        go.cell(go.max_row, 8).number_format = "0%"
    go.append(["Toplam", None, sum(len(g["satirlar"]) for g in s["plan"]), round(sum(g["km"] for g in s["plan"]), 1)])
    go.cell(go.max_row, 1).font = Font(bold=True)
    go.append([])
    go.append([f"Mesafe: kuş uçuşu × {p['katsayi']:g}; ortalama hız {p['hiz']:g} km/s; günlük süre {p['gunluk_dk']:g} dk. Başlangıç ve bitiş: "
               f"{s['baslangic'][0]:.4f}, {s['baslangic'][1]:.4f}. Süreler tahminidir."])

    sg = wb.create_sheet("Sığmayanlar")
    _baslik(sg, ["Müşteri No", "Müşteri", "İlçe", "Öncelik", "Son Ziyaret", "Gecikme (gün)", "Neden"], (11, 28, 14, 8, 11, 9, 50))
    for m, neden in s["sigmayan"]:
        sg.append([m.no, m.ad, m.ilce, m.oncelik, m.son, max(m.gecikme, 0), neden])
        sg.cell(sg.max_row, 5).number_format = "DD.MM.YYYY"
    for m in s["konumsuz"]:
        sg.append([m.no, m.ad, m.ilce, m.oncelik, m.son, max(m.gecikme, 0), "Konum (enlem / boylam) yok"])
        sg.cell(sg.max_row, 5).number_format = "DD.MM.YYYY"

    gr = wb.create_sheet("Bu Hafta Gerekmeyenler")
    _baslik(gr, ["Müşteri No", "Müşteri", "İlçe", "Öncelik", "Sıklık (gün)", "Son Ziyaret", "Sonraki Ziyaret"], (11, 28, 14, 8, 9, 11, 12))
    for m in s["gerekmeyen"]:
        gr.append([m.no, m.ad, m.ilce, m.oncelik, m.siklik, m.son, m.vade])
        gr.cell(gr.max_row, 6).number_format = gr.cell(gr.max_row, 7).number_format = "DD.MM.YYYY"

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Müşteri", "Açıklama"], (9, 20, 34, 90))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["kim"], u["aciklama"]])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(yol: Path, cikti: Path, hafta: date, baslangic: tuple[float, float], *, varsayilan_sure: float = 30, **kw) -> dict:
    musteriler, uy = oku(yol, varsayilan_sure)
    if not musteriler:
        raise ValueError(f"{yol.name}: müşteri bulunamadı")
    s = planla(musteriler, hafta, baslangic, **kw)
    s["uyarilar"] = s["uyarilar"] + uy
    rapor_yaz(cikti, s)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Müşteri konumu ve ziyaret sıklığına göre haftalık ziyaret planı ve rota sırası çıkarır.")
    p.add_argument("--musteriler", type=Path, default=ORNEK / "musteriler.csv",
                   help="Müşteri No, Müşteri, İlçe, Enlem, Boylam, Ziyaret Sıklığı, Ziyaret Süresi (dk), Öncelik (A/B/C), Son Ziyaret, Uygun Günler")
    p.add_argument("--hafta", help="Haftanın ilk günü GG.AA.YYYY (örnek veride 12.10.2026)")
    p.add_argument("--baslangic", help="Başlangıç / bitiş noktası 'enlem,boylam' (örnek veride 40.99,29.12)")
    p.add_argument("--gun-sayisi", type=int, default=5, help="Haftada kaç gün ziyaret (varsayılan 5; en çok 6)")
    p.add_argument("--gunluk-dk", type=float, default=480, help="Günlük toplam süre: yol + ziyaret (dk; varsayılan 480)")
    p.add_argument("--hiz", type=float, default=30, help="Ortalama hız km/s (varsayılan 30, şehir içi)")
    p.add_argument("--yol-katsayisi", type=float, default=1.3, help="Kuş uçuşu mesafeyi yol mesafesine çeviren katsayı (varsayılan 1,3)")
    p.add_argument("--ziyaret-dk", type=float, default=30, help="Ziyaret süresi boşsa (dk; varsayılan 30)")
    p.add_argument("--baslama", default="09:00", help="Güne başlama saati (varsayılan 09:00)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "ziyaret_plani.xlsx")
    a = p.parse_args(argv)
    if not a.musteriler.exists():
        print(f"[X] Dosya bulunamadı: {a.musteriler}")
        return 1
    ornek = a.musteriler == ORNEK / "musteriler.csv"
    hafta = tarih(a.hafta) if a.hafta else (date(2026, 10, 12) if ornek else date.today() - timedelta(days=date.today().weekday()) + timedelta(days=7))
    bas = None
    if a.baslangic or ornek:
        try:
            la, lo = (float(v) for v in (a.baslangic or "40.99,29.12").replace(";", ",").split(","))
            bas = (la, lo)
        except ValueError:
            bas = None
    try:
        datetime.strptime(a.baslama, "%H:%M")
        saat_ok = True
    except ValueError:
        saat_ok = False
    if not hafta or not bas or not saat_ok or not 1 <= a.gun_sayisi <= 6 or a.hiz <= 0:
        print("[X] --hafta GG.AA.YYYY, --baslangic 'enlem,boylam' (zorunlu), --baslama SS:DD, --gun-sayisi 1–6 olmalı")
        return 2
    try:
        s = calistir(a.musteriler, a.cikti, hafta, bas, varsayilan_sure=a.ziyaret_dk, gun_sayisi=a.gun_sayisi, gunluk_dk=a.gunluk_dk, hiz=a.hiz,
                     katsayi=a.yol_katsayisi, baslama=a.baslama)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] Hafta {hafta:%d.%m.%Y}: ziyareti gelen {s['gelen']} müşteri · planlanan {sum(len(g['satirlar']) for g in s['plan'])} · "
          f"sığmayan {len(s['sigmayan'])} · toplam yol {sum(g['km'] for g in s['plan']):.0f} km")
    for g in s["plan"]:
        print(f"     {GUNLER[g['gun']]:<10} {len(g['satirlar']):>2} ziyaret · {g['km']:5.1f} km · {g['dk']:4.0f} dk")
    for u in s["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['kim']} · {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
