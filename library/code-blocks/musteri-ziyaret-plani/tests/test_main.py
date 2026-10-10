import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
HAFTA = date(2026, 10, 12)
BAS = (40.99, 29.12)


def m(no, la, lo, **kw):
    return main.Musteri(no, no, "", (la, lo) if la is not None else None, kw.get("siklik", 7), kw.get("sure", 30), kw.get("oncelik", "B"), kw.get("son"),
                        kw.get("gunler", set()))


class GeometriTesti(unittest.TestCase):
    def test_haversine(self):
        self.assertAlmostEqual(main.haversine((41.0, 29.0), (42.0, 29.0)), 111.2, delta=0.2)      # 1 enlem derecesi ≈ 111,2 km
        self.assertEqual(main.haversine(BAS, BAS), 0)

    def test_rota_2opt(self):
        # Doğu yönünde bir hat üzerinde 4 nokta: en iyi tur sırayla gidip dönmektir
        nokta = [m("D", 40.99, 29.20), m("B", 40.99, 29.14), m("C", 40.99, 29.17), m("A", 40.99, 29.13)]
        sira = [x.no for x in main.rota_sirala(BAS, nokta)]
        self.assertIn(sira, (["A", "B", "C", "D"], ["D", "C", "B", "A"]))
        km = main.rota_km(BAS, [x.konum for x in main.rota_sirala(BAS, nokta)], 1.0)
        self.assertAlmostEqual(km, 2 * main.haversine(BAS, (40.99, 29.20)), places=6)

    def test_siklik_ve_gunler(self):
        self.assertEqual([main.siklik_gun(x) for x in ("Haftalık", "2 haftada bir", "Aylık", "3 ayda bir", "10 gün", "ara sıra")], [7, 14, 28, 84, 10, None])
        self.assertEqual(main.gunleri_coz("Sal, Per"), {1, 3})
        self.assertEqual(main.gunleri_coz("Pazartesi/Cuma"), {0, 4})


class PlanTesti(unittest.TestCase):
    def test_vade_kisit_ve_kapasite(self):
        lst = [m("A1", 41.00, 29.13, son=date(2026, 10, 6), oncelik="A"),                 # vade 13.10 → gelir
               m("B1", 41.00, 29.14, son=date(2026, 10, 10), siklik=14),                   # vade 24.10 → bu hafta gerekmez
               m("K1", 40.98, 29.10, gunler={3}),                                          # yalnız Perşembe; hiç ziyaret edilmemiş
               m("X1", 40.98, 29.11, gunler={5}),                                          # yalnız Cumartesi: 5 günlük planın dışında
               m("N1", None, None)]
        s = main.planla(lst, HAFTA, BAS)
        planli = {x["m"].no: g["gun"] for g in s["plan"] for x in g["satirlar"]}
        self.assertEqual(set(planli), {"A1", "K1"})
        self.assertEqual(planli["K1"], 3)
        self.assertEqual([x.no for x in s["gerekmeyen"]], ["B1"])
        self.assertEqual([x.no for x in s["konumsuz"]], ["N1"])
        self.assertEqual([(x.no, n.split(";")[0]) for x, n in s["sigmayan"]], [("X1", "Uygun günleri plan günlerinin dışında")])

    def test_sure_asimi(self):
        lst = [m(f"M{i}", 40.99 + i * 0.002, 29.13, sure=60, oncelik="C" if i == 2 else "A") for i in range(4)]
        s = main.planla(lst, HAFTA, BAS, gun_sayisi=1, gunluk_dk=200)
        self.assertEqual(len(s["plan"][0]["satirlar"]), 3)
        self.assertEqual(s["sigmayan"][0][0].no, "M2")                    # önceliği en düşük olan çıkarılır
        self.assertLessEqual(s["plan"][0]["dk"], 200)


class OrnekVeriTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "z.xlsx"
            s = main.calistir(ORNEK / "musteriler.csv", cikti, HAFTA, BAS)
            planli = [x["m"].no for g in s["plan"] for x in g["satirlar"]]
            self.assertEqual(len(planli), len(set(planli)))                             # kimse iki kez planlanmaz
            self.assertEqual(len(planli) + len(s["sigmayan"]) + len(s["konumsuz"]), s["gelen"])
            self.assertEqual(len(planli) + len(s["sigmayan"]) + len(s["konumsuz"]) + len(s["gerekmeyen"]), 48)
            for g in s["plan"]:
                self.assertLessEqual(g["dk"], 480)
                self.assertEqual([x["sira"] for x in g["satirlar"]], list(range(1, len(g["satirlar"]) + 1)))
            self.assertIn("M012", planli)                                                # hiç ziyaret edilmemiş
            turler = {(u["tur"], u["kim"].split()[0]) for u in s["uyarilar"]}
            self.assertTrue({("Konum yok", "M009"), ("Ziyaret gecikmiş", "M021"), ("Sıklık okunamadı", "M031")} <= turler)
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Haftalık Plan", "Gün Özeti", "Sığmayanlar", "Bu Hafta Gerekmeyenler", "Uyarılar"])
            self.assertTrue(wb["Haftalık Plan"]["M2"].hyperlink.target.startswith("https://www.google.com/maps/search/"))

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--baslangic", "kadıköy"]), 2)
            self.assertEqual(main.main(["--musteriler", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
