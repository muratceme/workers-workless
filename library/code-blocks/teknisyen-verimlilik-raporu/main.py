"""
Teknisyen Verimlilik Raporu — Workers / Workless kod bloğu
Satış Sonrası Hizmetler (Servis) › Satış Sonrası Hizmetler Müdürü
İş emirlerinden teknisyen bazında satılan saat, çalışılan saat ve mevcut (puantaj) saatini toplayıp atölye oranlarını hesaplar.
  Verimlilik %   = satılan (faturalanan / standart) saat ÷ çalışılan (işe kayıtlı) saat × 100
  Kullanım %     = çalışılan saat ÷ mevcut (işyerinde bulunulan) saat × 100
  Üretkenlik %   = satılan saat ÷ mevcut saat × 100  (= verimlilik × kullanım ÷ 100)
  Boş saat       = mevcut − çalışılan (işe kayıtlı olmayan süre)
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
IS = {"no": ("is emri no", "is emri", "no"), "tarih": ("tarih", "is emri tarihi", "kapanis tarihi"), "teknisyen": ("teknisyen", "usta", "personel"),
      "tur": ("is turu", "odeme turu", "tur", "hesap turu"), "satilan": ("satilan saat", "faturalanan saat", "standart saat", "flat rate"),
      "calisilan": ("calisilan saat", "fiili saat", "kayitli saat", "gerceklesen saat")}
DEVAM = {"teknisyen": ("teknisyen", "usta", "personel"), "tarih": ("tarih", "gun"), "mevcut": ("mevcut saat", "calisma saati", "puantaj saati", "saat")}
SIFIR = Decimal(0)


def _oran(a: Decimal, b: Decimal) -> Decimal | None:
    return yuvarla(a * 100 / b) if b else None


def is_oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, IS, ("teknisyen", "satilan", "calisilan")):
        if not metin(r["teknisyen"]) and not metin(r["no"]):
            continue
        kim = metin(r["no"]) or f"Satır {r['_satir']}"
        g = tarih(r["tarih"]) if metin(r["tarih"]) else None
        s, c = para(r["satilan"]) if metin(r["satilan"]) else SIFIR, para(r["calisilan"]) if metin(r["calisilan"]) else SIFIR
        if not metin(r["teknisyen"]) or s is None or c is None or s < 0 or c < 0 or (metin(r["tarih"]) and g is None):
            uy.append(uyari("Yüksek", "Satır okunamadı", kim, f"Satır {r['_satir']}: teknisyen, tarih ya da saatler okunamadı."))
            continue
        l.append({"no": metin(r["no"]), "tarih": g, "teknisyen": metin(r["teknisyen"]), "tur": metin(r["tur"]) or "Belirtilmemiş", "satilan": s, "calisilan": c})
    return l, uy


def devam_oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, DEVAM, ("teknisyen", "mevcut")):
        if not metin(r["teknisyen"]):
            continue
        g = tarih(r["tarih"]) if metin(r["tarih"]) else None
        m = para(r["mevcut"])
        if m is None or m < 0 or m > 24 * 31 or (metin(r["tarih"]) and g is None):
            uy.append(uyari("Yüksek", "Puantaj satırı okunamadı", metin(r["teknisyen"]), f"Satır {r['_satir']}: tarih ya da mevcut saat okunamadı."))
            continue
        l.append({"teknisyen": metin(r["teknisyen"]), "tarih": g, "mevcut": m})
    return l, uy


def analiz(isler: list[dict], devam: list[dict], verim_alt: Decimal = Decimal(85), kullanim_alt: Decimal = Decimal(80)) -> dict:
    if not isler:
        raise ValueError("Okunabilen iş emri satırı yok.")
    uy = []
    mevcut = {}
    for d in devam:
        mevcut[katla(d["teknisyen"])] = mevcut.get(katla(d["teknisyen"]), SIFIR) + d["mevcut"]
    gunluk_mevcut = {}
    for d in devam:
        if d["tarih"]:
            k = (katla(d["teknisyen"]), d["tarih"])
            gunluk_mevcut[k] = gunluk_mevcut.get(k, SIFIR) + d["mevcut"]
    gunluk_calisma = {}
    for i in isler:
        if i["tarih"]:
            k = (katla(i["teknisyen"]), i["tarih"])
            gunluk_calisma[k] = gunluk_calisma.get(k, SIFIR) + i["calisilan"]
    for (t, g), c in sorted(gunluk_calisma.items()):
        m = gunluk_mevcut.get((t, g))
        if m is not None and c > m:
            ad = next(i["teknisyen"] for i in isler if katla(i["teknisyen"]) == t)
            uy.append(uyari("Yüksek", "Çalışılan saat mevcut saati aşıyor", f"{ad} · {g:%d.%m.%Y}", f"İş emirlerine {c} saat yazılmış, puantajda {m} saat var; süre kaydını kontrol edin."))
    teknisyen = []
    adlar = {}
    for i in isler:
        adlar.setdefault(katla(i["teknisyen"]), i["teknisyen"])
    for k, ad in adlar.items():
        l = [i for i in isler if katla(i["teknisyen"]) == k]
        s = sum((i["satilan"] for i in l), SIFIR)
        c = sum((i["calisilan"] for i in l), SIFIR)
        m = mevcut.get(k)
        v, ku, u = _oran(s, c), _oran(c, m) if m is not None else None, _oran(s, m) if m is not None else None
        teknisyen.append({"ad": ad, "is": len({i["no"] for i in l if i["no"]}) or len(l), "satilan": s, "calisilan": c, "mevcut": m, "bos": m - c if m is not None else None,
                          "verimlilik": v, "kullanim": ku, "uretkenlik": u})
        if m is None and devam:
            uy.append(uyari("Orta", "Puantaj yok", ad, "Teknisyenin puantaj (mevcut saat) kaydı yok; kullanım ve üretkenlik hesaplanmadı."))
        if v is not None and v < verim_alt:
            uy.append(uyari("Bilgi", "Verimlilik düşük", ad, f"Verimlilik %{v} (eşik %{verim_alt}); iş dağılımı, eğitim, parça bekleme ve tekrar işleri değerlendirin."))
        if v is not None and v > 150:
            uy.append(uyari("Bilgi", "Verimlilik çok yüksek", ad, f"Verimlilik %{v}; çalışılan saatlerin eksik kaydedilmediğini kontrol edin."))
        if ku is not None and ku < kullanim_alt:
            uy.append(uyari("Bilgi", "Kullanım düşük", ad, f"Kullanım %{ku} (eşik %{kullanim_alt}); {m - c} saat işe kayıtlı değil (iş yükü, randevu planı ya da kayıt alışkanlığı)."))
    for d in {katla(d["teknisyen"]): d["teknisyen"] for d in devam}.items():
        if d[0] not in adlar:
            uy.append(uyari("Bilgi", "İş emri olmayan teknisyen", d[1], f"Puantajda {mevcut[d[0]]} saat var ama iş emri satırı yok."))
    teknisyen.sort(key=lambda s: (s["uretkenlik"] is None, -(s["uretkenlik"] or 0), -(s["verimlilik"] or 0), katla(s["ad"])))
    turler = []
    for tur in sorted({i["tur"] for i in isler}, key=katla):
        l = [i for i in isler if i["tur"] == tur]
        s, c = sum((i["satilan"] for i in l), SIFIR), sum((i["calisilan"] for i in l), SIFIR)
        turler.append({"ad": tur, "is": len(l), "satilan": s, "calisilan": c, "verimlilik": _oran(s, c)})
    s = sum((t["satilan"] for t in teknisyen), SIFIR)
    c = sum((t["calisilan"] for t in teknisyen), SIFIR)
    mt = [t for t in teknisyen if t["mevcut"] is not None]
    m = sum((t["mevcut"] for t in mt), SIFIR) if mt else None
    cm = sum((t["calisilan"] for t in mt), SIFIR)
    sm = sum((t["satilan"] for t in mt), SIFIR)
    toplam = {"ad": "ATÖLYE", "is": sum(t["is"] for t in teknisyen), "satilan": s, "calisilan": c, "mevcut": m, "bos": m - cm if m is not None else None, "verimlilik": _oran(s, c),
              "kullanim": _oran(cm, m) if m else None, "uretkenlik": _oran(sm, m) if m else None}
    return {"teknisyen": teknisyen, "turler": turler, "toplam": toplam, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, verim_alt: Decimal, kullanim_alt: Decimal) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Teknisyen"
    _baslik(ws, ["Teknisyen", "İş Emri", "Satılan Saat", "Çalışılan Saat", "Mevcut Saat", "Boş Saat", "Verimlilik %", "Kullanım %", "Üretkenlik %"], (24, 9, 11, 11, 11, 9, 11, 11, 11))
    for s in r["teknisyen"] + [r["toplam"]]:
        n = _satir(ws, [s["ad"], s["is"], s["satilan"], s["calisilan"], s["mevcut"], s["bos"], s["verimlilik"], s["kullanim"], s["uretkenlik"]], para_s=(3, 4, 5, 6, 7, 8, 9),
                   dolgu={7: SARI if s["verimlilik"] is not None and s["verimlilik"] < verim_alt else None, 8: SARI if s["kullanim"] is not None and s["kullanim"] < kullanim_alt else None})
        if s["ad"] == "ATÖLYE":
            for h in ws[n]:
                h.font = Font(bold=True)
    ws.append([])
    ws.append(["Verimlilik = satılan ÷ çalışılan · Kullanım = çalışılan ÷ mevcut · Üretkenlik = satılan ÷ mevcut. Atölye kullanım ve üretkenliği yalnız puantajı olan teknisyenlerden hesaplanır."])
    ws = wb.create_sheet("İş Türü")
    _baslik(ws, ["İş Türü", "Satır", "Satılan Saat", "Çalışılan Saat", "Verimlilik %"], (20, 8, 11, 11, 11))
    for s in r["turler"]:
        _satir(ws, [s["ad"], s["is"], s["satilan"], s["calisilan"], s["verimlilik"]], para_s=(3, 4, 5))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(is_yol: Path, devam_yol: Path | None, cikti: Path, verim_alt: Decimal = Decimal(85), kullanim_alt: Decimal = Decimal(80)) -> dict:
    isler, uy1 = is_oku(is_yol)
    devam, uy2 = devam_oku(devam_yol) if devam_yol else ([], [])
    r = analiz(isler, devam, verim_alt, kullanim_alt)
    r["uyarilar"] = sirala(uy1 + uy2 + r["uyarilar"])
    rapor_yaz(cikti, r, verim_alt, kullanim_alt)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Teknisyen bazında satılan / çalışılan / mevcut saat ve verimlilik, kullanım, üretkenlik oranlarını hesaplar.")
    p.add_argument("--is-emirleri", type=Path, help="İş Emri No, Tarih, Teknisyen, İş Türü (Müşteri / Garanti / İç), Satılan Saat, Çalışılan Saat")
    p.add_argument("--puantaj", type=Path, help="Teknisyen, Tarih, Mevcut Saat (isteğe bağlı; verilmezse kullanım ve üretkenlik hesaplanmaz)")
    p.add_argument("--verim-alt", default="85", help="Verimlilik bu yüzdenin altındaysa bilgi notu (varsayılan 85)")
    p.add_argument("--kullanim-alt", default="80", help="Kullanım bu yüzdenin altındaysa bilgi notu (varsayılan 80)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "teknisyen_verimlilik.xlsx")
    a = p.parse_args(argv)
    is_yol = a.is_emirleri or ORNEK / "is_emirleri.csv"
    devam = a.puantaj or (ORNEK / "puantaj.csv" if a.is_emirleri is None else None)
    va, ka = para(a.verim_alt), para(a.kullanim_alt)
    if va is None or not 0 <= va <= 300 or ka is None or not 0 <= ka <= 100:
        print("[X] --verim-alt 0-300, --kullanim-alt 0-100 olmalı")
        return 2
    for y in (is_yol, devam):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(is_yol, devam, a.cikti, va, ka)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] Atölye: satılan {t['satilan']} · çalışılan {t['calisilan']} · mevcut {t['mevcut']} saat · verimlilik %{t['verimlilik']} · kullanım %{t['kullanim']} · üretkenlik %{t['uretkenlik']}")
    for s in r["teknisyen"]:
        print(f"     {s['ad']}: verimlilik %{s['verimlilik']} · kullanım %{s['kullanim']} · üretkenlik %{s['uretkenlik']}")
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
