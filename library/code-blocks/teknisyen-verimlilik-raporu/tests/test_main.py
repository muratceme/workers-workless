import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def i(no, gun, tek, s, c, tur="Müşteri"):
    return {"no": no, "tarih": date(2026, 9, gun), "teknisyen": tek, "tur": tur, "satilan": D(s), "calisilan": D(c)}


def d(tek, gun, m):
    return {"teknisyen": tek, "tarih": date(2026, 9, gun), "mevcut": D(m)}


class AnalizTesti(unittest.TestCase):
    def test_oranlar(self):
        isler = [i("1", 1, "A", 4, 3), i("2", 1, "A", 3, 3, "Garanti"), i("3", 2, "A", 2, 2), i("4", 1, "B", 2, 4), i("5", 1, "B", 1, 1, "Garanti"), i("6", 1, "C", 9, 9)]
        devam = [d("A", 1, "7.5"), d("A", 2, "7.5"), d("B", 1, 8), d("Z", 1, 8), d("C", 1, 8)]
        r = main.analiz(isler, devam)
        x = {s["ad"]: s for s in r["teknisyen"]}
        # A: satılan 9 · çalışılan 8 · mevcut 15 → verimlilik %112,5, kullanım %53,33, üretkenlik %60
        self.assertEqual([x["A"][k] for k in ("is", "satilan", "calisilan", "mevcut", "bos", "verimlilik", "kullanim", "uretkenlik")],
                         [3, D(9), D(8), D(15), D(7), D("112.50"), D("53.33"), D("60.00")])
        self.assertEqual([x["B"][k] for k in ("verimlilik", "kullanim", "uretkenlik")], [D("60.00"), D("62.50"), D("37.50")])
        self.assertEqual([s["ad"] for s in r["teknisyen"]], ["C", "A", "B"])
        self.assertEqual([(t["ad"], t["satilan"], t["calisilan"], t["verimlilik"]) for t in r["turler"]], [("Garanti", D(4), D(4), D(100)), ("Müşteri", D(17), D(18), D("94.44"))])
        t = r["toplam"]
        self.assertEqual((t["satilan"], t["calisilan"], t["mevcut"], t["verimlilik"], t["kullanim"], t["uretkenlik"]), (D(21), D(22), D(31), D("95.45"), D("70.97"), D("67.74")))
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Bilgi", "Kullanım düşük", "A"), ("Bilgi", "Kullanım düşük", "B"), ("Bilgi", "Verimlilik düşük", "B"), ("Bilgi", "İş emri olmayan teknisyen", "Z"),
                          ("Yüksek", "Çalışılan saat mevcut saati aşıyor", "C · 01.09.2026")])
        r = main.analiz([i("1", 1, "A", 1, 1)], [d("B", 1, 8)])
        self.assertIn(("Orta", "Puantaj yok"), [(u["onem"], u["tur"]) for u in r["uyarilar"]])
        r = main.analiz([i("1", 1, "A", 1, "0.5")], [])
        self.assertEqual((r["toplam"]["kullanim"], r["teknisyen"][0]["verimlilik"], r["uyarilar"][0]["tur"]), (None, D(200), "Verimlilik çok yüksek"))
        with self.assertRaises(ValueError):
            main.analiz([], [])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "is_emirleri.csv", ORNEK / "puantaj.csv", cikti)
            t = r["toplam"]
            self.assertEqual((t["satilan"], t["calisilan"], t["mevcut"], t["verimlilik"], t["kullanim"], t["uretkenlik"]),
                             (D("879.6"), D("921.4"), D("1140.0"), D("95.46"), D("80.82"), D("77.16")))
            self.assertEqual([s["ad"] for s in r["teknisyen"]][:2], ["Usta 1 (Kıdemli)", "Usta 2"])
            self.assertEqual([u["kim"] for u in r["uyarilar"] if u["onem"] == "Yüksek"], ["Usta 2 · 17.09.2026"])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Teknisyen", "İş Türü", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            iy, py = Path(tmp) / "i.csv", Path(tmp) / "p.csv"
            iy.write_text("İş Emri;Tarih;Usta;Faturalanan Saat;Fiili Saat\n1;01.09.2026;A;1,5;1,2\n2;;;1;1\n3;x;A;1;1\n4;01.09.2026;A;-1;1\n5;01.09.2026;B;;2\n", encoding="utf-8")
            py.write_text("Usta;Gün;Puantaj Saati\nA;01.09.2026;7,5\nB;x;8\nB;;800\n", encoding="utf-8")
            l, uy = main.is_oku(iy)
            self.assertEqual(([(x["no"], x["satilan"], x["calisilan"]) for x in l], len(uy)), ([("1", D("1.5"), D("1.2")), ("5", D(0), D(2))], 3))
            l, uy = main.devam_oku(py)
            self.assertEqual((len(l), len(uy)), (1, 2))
            c = str(Path(tmp) / "x.xlsx")
            self.assertEqual(main.main(["--is-emirleri", str(iy), "--puantaj", str(py), "--verim-alt", "90", "--kullanim-alt", "75", "--cikti", c]), 0)
            self.assertEqual(main.main(["--is-emirleri", str(iy), "--cikti", c]), 0)
            self.assertEqual([main.main(["--cikti", c] + v) for v in (["--verim-alt", "x"], ["--kullanim-alt", "101"])], [2, 2])
            self.assertEqual(main.main(["--is-emirleri", str(Path(tmp) / "yok.csv")]), 1)
            iy.write_text("Usta;Faturalanan Saat;Fiili Saat\n;;\n", encoding="utf-8")
            self.assertEqual(main.main(["--is-emirleri", str(iy), "--cikti", c]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
