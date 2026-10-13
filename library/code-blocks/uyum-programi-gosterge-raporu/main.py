"""
Uyum Programı Gösterge Raporu — Workers / Workless kod bloğu
Kurumsal Uyum › Uyum Müdürü

İnceleme, bildirim, eğitim ve yaptırım taraması göstergelerini dönemsel uyum raporunda toplar.
  Son dönem          = dosyadaki en büyük dönem etiketi (dönemler metin olarak sıralanır: 2026-Q1 < 2026-Q2)
  Değişim %          = (son dönem − önceki dönem) ÷ |önceki dönem| × 100
  Eğilim             = yönü "Yüksek iyi" olan göstergede artış İyileşti, azalış Kötüleşti; "Düşük iyi" olanda tersi; yönsüzde Arttı / Azaldı
  Hedef durumu       = Yüksek iyi: değer ≥ hedef → Hedefte · Düşük iyi: değer ≤ hedef → Hedefte · aksi hâlde Hedef dışı
  Hedefe sapma %     = (değer − hedef) ÷ |hedef| × 100
  Belirgin kötüleşme = kötü yöndeki değişim --degisim (25) yüzde ve üstünde
Göstergeleri, hedefleri ve yönlerini siz tanımlarsınız. Rapor göstergeleri toplar ve karşılaştırır; uyum programının yeterliliği hakkında hüküm vermez.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
SUTUN = {"donem": ("donem", "ay", "ceyrek", "rapor donemi"), "ad": ("gosterge", "gosterge adi", "olcut"), "deger": ("deger", "gerceklesen", "sonuc"), "hedef": ("hedef", "hedef deger"), "yon": ("yon", "iyi yon", "hedef yonu"),
         "birim": ("birim", "olcu birimi")}
YON = {"yuksek iyi": 1, "yuksek": 1, "artmali": 1, "yukari": 1, "en az": 1, "dusuk iyi": -1, "dusuk": -1, "azalmali": -1, "asagi": -1, "en cok": -1, "": 0, "yok": 0, "izleme": 0}
YON_AD = {1: "Yüksek iyi", -1: "Düşük iyi", 0: ""}
DURUM_RENK = {"Hedef dışı": KIRMIZI, "Hedefte": YESIL, "Hedef yok": None, "Veri yok": MAVI}
YUZ = Decimal(100)


def olcu(x) -> Decimal | None:
    """Değer okur: virgül ya da tek nokta ondalık ayırıcıdır (4,2 = 4.2); '1.250,5' binlik ayırıcılıdır."""
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x or ""))
    try:
        return Decimal(s.replace(".", "").replace(",", ".") if "," in s else s) if s else None
    except InvalidOperation:
        return None


def donem_etiketi(x) -> str:
    if isinstance(x, (date, datetime)):
        return f"{x:%Y-%m}"
    return metin(x).upper().replace("Ç", "Q").replace("/", "-").replace(" ", "")


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("donem", "ad", "deger")):
        ad, donem = metin(r["ad"]), donem_etiketi(r["donem"])
        if not ad and not donem:
            continue
        deger, hedef, yon = olcu(r["deger"]), olcu(r["hedef"]), YON.get(katla(r["yon"]))
        if not ad or not donem or deger is None or yon is None or (metin(r["hedef"]) and hedef is None):
            uy.append(uyari("Yüksek", "Satır okunamadı", ad or f"Satır {r['_satir']}", f"Satır {r['_satir']}: dönem '{metin(r['donem'])}', değer '{metin(r['deger'])}', hedef '{metin(r['hedef'])}', yön '{metin(r['yon'])}' "
                            "(Yüksek iyi / Düşük iyi ya da boş). Satır alınmadı."))
            continue
        l.append({"donem": donem, "ad": ad, "deger": deger, "hedef": hedef, "yon": yon, "birim": metin(r["birim"])})
    return l, uy


def _oran(fark: Decimal, taban: Decimal | None) -> Decimal | None:
    return yuvarla(fark * YUZ / abs(taban)) if taban else None


def raporla(kayit: list[dict], esik: Decimal = Decimal(25), donem_sayisi: int = 4) -> dict:
    if not kayit:
        raise ValueError("Raporlanacak gösterge yok.")
    uy = []
    donemler = sorted({k["donem"] for k in kayit})
    son, gosterilen = donemler[-1], donemler[-donem_sayisi:]
    grup = {}
    for k in kayit:
        g = grup.setdefault(katla(k["ad"]), {"ad": k["ad"], "seri": {}, "satirlar": []})
        if k["donem"] in g["seri"]:
            uy.append(uyari("Orta", "Aynı dönem iki kez girilmiş", k["ad"], f"{k['donem']} dönemi için birden çok değer var; son satırdaki değer kullanıldı."))
        g["seri"][k["donem"]] = k["deger"]
        g["satirlar"].append(k)
    gostergeler = []
    for g in grup.values():
        en_son = max(g["satirlar"], key=lambda k: (k["donem"], g["satirlar"].index(k)))
        sirali = sorted(g["seri"])
        o = {"ad": g["ad"], "birim": en_son["birim"], "yon": en_son["yon"], "seri": [g["seri"].get(d) for d in gosterilen], "son_donem": sirali[-1], "deger": g["seri"][sirali[-1]], "onceki": g["seri"][sirali[-2]] if len(sirali) > 1 else None,
             "hedef": en_son["hedef"]}
        o["degisim"] = _oran(o["deger"] - o["onceki"], o["onceki"]) if o["onceki"] is not None else None
        fark = o["deger"] - o["onceki"] if o["onceki"] is not None else None
        o["egilim"] = None if fark is None else "Değişmedi" if fark == 0 else ("İyileşti" if fark * o["yon"] > 0 else "Kötüleşti") if o["yon"] else ("Arttı" if fark > 0 else "Azaldı")
        o["sapma"] = _oran(o["deger"] - o["hedef"], o["hedef"]) if o["hedef"] is not None else None
        if o["son_donem"] != son:
            o["durum"] = "Veri yok"
            uy.append(uyari("Orta", "Son dönemde veri yok", o["ad"], f"{son} dönemi için değer girilmemiş; son değer {o['son_donem']} dönemine ait ({o['deger']:f})."))
        elif o["hedef"] is None or not o["yon"]:
            o["durum"] = "Hedef yok"
            if o["hedef"] is not None:
                uy.append(uyari("Bilgi", "Hedef var, yön yok", o["ad"], "Hedef yazılmış ama göstergenin iyi yönü (Yüksek iyi / Düşük iyi) belirtilmemiş; hedef karşılaştırması yapılmadı."))
        else:
            o["durum"] = "Hedefte" if (o["deger"] - o["hedef"]) * o["yon"] >= 0 else "Hedef dışı"
            if o["durum"] == "Hedef dışı":
                uy.append(uyari("Yüksek", "Hedef dışı", o["ad"], f"{son}: değer {o['deger']:f}, hedef {'en az' if o['yon'] > 0 else 'en çok'} {o['hedef']:f}" + (f" (sapma %{o['sapma']})." if o["sapma"] is not None else ".")))
        if o["son_donem"] == son and o["degisim"] is not None and abs(o["degisim"]) >= esik:
            if o["egilim"] == "Kötüleşti":
                uy.append(uyari("Orta", "Belirgin kötüleşme", o["ad"], f"Önceki döneme göre %{abs(o['degisim'])} kötüleşti ({o['onceki']:f} → {o['deger']:f})."))
            elif not o["yon"]:
                uy.append(uyari("Bilgi", "Belirgin değişim", o["ad"], f"Önceki döneme göre %{o['degisim']} değişti ({o['onceki']:f} → {o['deger']:f}). Nedenini raporda açıklayın."))
        gostergeler.append(o)
    return {"gostergeler": gostergeler, "donemler": gosterilen, "son": son, "say": {d: sum(1 for o in gostergeler if o["durum"] == d) for d in DURUM_RENK}, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Gösterge Tablosu"
    d = r["donemler"]
    _baslik(ws, ["Gösterge", "Birim", "İyi Yön"] + d + ["Değişim %", "Eğilim", "Hedef", "Hedefe Sapma %", "Durum", "Açıklama (siz doldurun)"], (46, 9, 11) + (12,) * len(d) + (10, 11, 10, 11, 12, 40))
    n = len(d)
    for o in sorted(r["gostergeler"], key=lambda o: (list(DURUM_RENK).index(o["durum"]), katla(o["ad"]))):
        _satir(ws, [o["ad"], o["birim"], YON_AD[o["yon"]]] + o["seri"] + [o["degisim"], o["egilim"], o["hedef"], o["sapma"], o["durum"], None], para_s=tuple(range(4, 4 + n)) + (4 + n, 6 + n, 7 + n),
               dolgu={8 + n: DURUM_RENK[o["durum"]], 5 + n: {"Kötüleşti": KIRMIZI, "İyileşti": YESIL}.get(o["egilim"])}, sar=(1,))
    ws = wb.create_sheet("Özet")
    _baslik(ws, ["Durum", "Gösterge Sayısı"], (20, 14))
    for durum, sayi in r["say"].items():
        _satir(ws, [durum, sayi], dolgu={1: DURUM_RENK[durum]})
    _satir(ws, ["Toplam", len(r["gostergeler"])])
    ws.append([])
    ws.append([f"Rapor dönemi: {r['son']} · gösterilen dönemler: {', '.join(d)}"])
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, esik: Decimal = Decimal(25), donem_sayisi: int = 4) -> dict:
    kayit, uy = oku(yol)
    r = raporla(kayit, esik, donem_sayisi)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Uyum programı göstergelerini dönemsel olarak toplar; hedefle ve önceki dönemle karşılaştırır.")
    p.add_argument("--gostergeler", type=Path, help="Dönem, Gösterge, Değer, Hedef, Yön (Yüksek iyi / Düşük iyi), Birim")
    p.add_argument("--degisim", default="25", help="Önceki döneme göre değişim bu yüzde ve üstündeyse uyarılır (varsayılan 25)")
    p.add_argument("--donem", type=int, default=4, help="Tabloda gösterilecek son dönem sayısı (varsayılan 4)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "uyum_gosterge_raporu.xlsx")
    a = p.parse_args(argv)
    yol = a.gostergeler or ORNEK / "gostergeler.csv"
    esik = para(a.degisim)
    if esik is None or esik <= 0 or not 1 <= a.donem <= 24:
        print("[X] --degisim pozitif; --donem 1–24 arasında olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, esik, a.donem)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    s = r["say"]
    print(f"[OK] {r['son']} dönemi · {len(r['gostergeler'])} gösterge · hedefte {s['Hedefte']}, hedef dışı {s['Hedef dışı']}, hedefsiz {s['Hedef yok']}, verisi eksik {s['Veri yok']}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
