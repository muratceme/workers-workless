import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def k(ad, donem, deger, hedef=None, yon=0, birim=""):
    return {"donem": donem, "ad": ad, "deger": D(str(deger)), "hedef": None if hedef is None else D(str(hedef)), "yon": yon, "birim": birim}


def g(kayit, **kw):
    r = main.raporla(kayit, **kw)
    return {o["ad"]: (o["deger"], o["onceki"], o["degisim"], o["egilim"], o["sapma"], o["durum"]) for o in r["gostergeler"]}, sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"])


class RaporTesti(unittest.TestCase):
    def test_hedef_ve_egilim(self):
        l = [k("yüksek", "Q1", 90, 95, 1), k("yüksek", "Q2", 95, 95, 1), k("düşük", "Q1", 4, 5, -1), k("düşük", "Q2", 5, 5, -1), k("y_dışı", "Q1", 96, 95, 1), k("y_dışı", "Q2", "94.99", 95, 1), k("d_dışı", "Q1", "4.8", 5, -1),
             k("d_dışı", "Q2", "6.5", 5, -1), k("yönsüz", "Q1", 465), k("yönsüz", "Q2", 610), k("hedefsiz", "Q1", 10, None, 1), k("hedefsiz", "Q2", 10, None, 1), k("sıfır", "Q1", 0, 0, -1), k("sıfır", "Q2", 2, 0, -1),
             k("tek", "Q2", 7, 5, 1), k("yön_yok", "Q2", 7, 5)]
        s, uy = g(l)
        self.assertEqual(s, {"yüksek": (D(95), D(90), D("5.56"), "İyileşti", D("0.00"), "Hedefte"), "düşük": (D(5), D(4), D("25.00"), "Kötüleşti", D("0.00"), "Hedefte"),
                             "y_dışı": (D("94.99"), D(96), D("-1.05"), "Kötüleşti", D("-0.01"), "Hedef dışı"), "d_dışı": (D("6.5"), D("4.8"), D("35.42"), "Kötüleşti", D("30.00"), "Hedef dışı"),
                             "yönsüz": (D(610), D(465), D("31.18"), "Arttı", None, "Hedef yok"), "hedefsiz": (D(10), D(10), D("0.00"), "Değişmedi", None, "Hedef yok"), "sıfır": (D(2), D(0), None, "Kötüleşti", None, "Hedef dışı"),
                             "tek": (D(7), None, None, None, D("40.00"), "Hedefte"), "yön_yok": (D(7), None, None, None, D("40.00"), "Hedef yok")})
        self.assertEqual(uy, sorted([("Orta", "Belirgin kötüleşme", "düşük"), ("Yüksek", "Hedef dışı", "y_dışı"), ("Yüksek", "Hedef dışı", "d_dışı"), ("Orta", "Belirgin kötüleşme", "d_dışı"), ("Bilgi", "Belirgin değişim", "yönsüz"),
                                     ("Yüksek", "Hedef dışı", "sıfır"), ("Bilgi", "Hedef var, yön yok", "yön_yok")]))
        self.assertEqual([u[1] for u in g([k("a", "Q1", 100, 50, -1), k("a", "Q2", 76, 50, -1), k("b", "Q1", 100), k("b", "Q2", 76)])[1]], ["Hedef dışı"])
        self.assertEqual([u[1] for u in g([k("b", "Q1", 100), k("b", "Q2", 90)], esik=D(10))[1]], ["Belirgin değişim"])
        self.assertEqual(g([k("a", "Q1", 5), k("a", "Q2", 3)])[0]["a"][3], "Azaldı")
        with self.assertRaises(ValueError):
            main.raporla([])

    def test_donemler(self):
        l = [k("a", f"2026-Q{i}", i, 3, 1, "adet") for i in (3, 1, 2)] + [k("a", "2025-Q4", 9, 3, 1), k("eski", "2026-Q1", 70, 100, 1), k("eski", "2026-Q2", 85, 100, 1), k("çift", "2026-Q3", 1), k("ÇİFT", "2026-Q3", 2), k("a", "2025-Q3", 8)]
        r = main.raporla(l, donem_sayisi=3)
        o = {x["ad"]: x for x in r["gostergeler"]}
        self.assertEqual((r["donemler"], r["son"], o["a"]["seri"], o["a"]["deger"], o["a"]["onceki"], o["a"]["durum"], o["a"]["birim"], o["a"]["yon"]), (["2026-Q1", "2026-Q2", "2026-Q3"], "2026-Q3", [D(1), D(2), D(3)], D(3), D(2), "Hedefte", "adet", 1))
        self.assertEqual((o["eski"]["seri"], o["eski"]["son_donem"], o["eski"]["deger"], o["eski"]["durum"], o["çift"]["deger"], o["çift"]["seri"]), ([D(70), D(85), None], "2026-Q2", D(85), "Veri yok", D(2), [None, None, D(2)]))
        self.assertEqual((r["say"], sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"])),
                         ({"Hedef dışı": 0, "Hedefte": 1, "Hedef yok": 1, "Veri yok": 1}, [("Orta", "Aynı dönem iki kez girilmiş", "ÇİFT"), ("Orta", "Son dönemde veri yok", "eski")]))
        self.assertEqual(main.raporla(l)["donemler"], ["2025-Q4", "2026-Q1", "2026-Q2", "2026-Q3"])
        self.assertEqual([main.donem_etiketi(x) for x in ("2026-Ç3", "2026/q3", " 2026 q3 ", date(2026, 9, 30), "2026-09")], ["2026-Q3", "2026-Q3", "2026Q3", "2026-09", "2026-09"])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "gostergeler.csv", cikti)
            self.assertEqual((r["donemler"], r["son"], r["say"]), (["2026-Q1", "2026-Q2", "2026-Q3"], "2026-Q3", {"Hedef dışı": 2, "Hedefte": 3, "Hedef yok": 2, "Veri yok": 1}))
            self.assertEqual([(o["ad"], o["deger"], o["degisim"], o["egilim"], o["sapma"], o["durum"]) for o in r["gostergeler"]],
                             [("İncelenen uyarı sayısı", D(610), D("31.18"), "Arttı", None, "Hedef yok"), ("Ortalama uyarı inceleme süresi", D("6.5"), D("35.42"), "Kötüleşti", D("30.00"), "Hedef dışı"),
                              ("Şüpheli işlem bildirimi sayısı", D(14), D("-6.67"), "Azaldı", None, "Hedef yok"), ("Uyum eğitimini tamamlayan personel", D(96), D("5.49"), "İyileşti", D("1.05"), "Hedefte"),
                              ("Müşteri tanıma bilgisi eksik dosya oranı", D("1.8"), D("-25.00"), "İyileşti", D("-10.00"), "Hedefte"), ("Yaptırım taramasında bekleyen eşleşme", D(25), D("38.89"), "Kötüleşti", D("25.00"), "Hedef dışı"),
                              ("Süresi geçmiş denetim bulgusu", D(0), D("-100.00"), "İyileşti", None, "Hedefte"), ("Gözden geçirilen politika oranı", D(85), D("21.43"), "İyileşti", D("-15.00"), "Veri yok")])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Yüksek", "Hedef dışı", "Ortalama uyarı inceleme süresi"), ("Yüksek", "Hedef dışı", "Yaptırım taramasında bekleyen eşleşme"), ("Orta", "Belirgin kötüleşme", "Ortalama uyarı inceleme süresi"),
                              ("Orta", "Belirgin kötüleşme", "Yaptırım taramasında bekleyen eşleşme"), ("Orta", "Son dönemde veri yok", "Gözden geçirilen politika oranı"), ("Bilgi", "Belirgin değişim", "İncelenen uyarı sayısı")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Gösterge Tablosu", "Özet", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "g.csv"
            y.write_text("Rapor Dönemi;Ölçüt;Gerçekleşen;Hedef Değer;İyi Yön;Ölçü Birimi\n2026/ç3;Eğitim;96,5;95;YÜKSEK;%\n2026-Q3;Süre;1.250,5;;en çok;gün\n2026-Q3;X;;5;;\n2026-Q3;Y;5;x;;\n2026-Q3;Z;5;;sağa;\n;W;5;;;\n2026-Q3;;5;;;\n;;;;;\n",
                         encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual((l, [(u["tur"], u["kim"]) for u in uy]),
                             ([{"donem": "2026-Q3", "ad": "Eğitim", "deger": D("96.5"), "hedef": D(95), "yon": 1, "birim": "%"}, {"donem": "2026-Q3", "ad": "Süre", "deger": D("1250.5"), "hedef": None, "yon": -1, "birim": "gün"}],
                              [("Satır okunamadı", n) for n in ("X", "Y", "Z", "W", "Satır 8")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--gostergeler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--degisim", "20", "--donem", "6"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--degisim", "0"], ["--degisim", "x"], ["--donem", "0"], ["--donem", "25"])], [2] * 4)
            self.assertEqual(main.main(["--gostergeler", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Dönem;Gösterge\nA;B\n", "Dönem;Gösterge;Değer\n;;\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
