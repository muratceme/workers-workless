"""
Sektör Karşılaştırmalı Rasyo Analizi — Workers / Workless kod bloğu
Krediler Tahsis › Kredi Tahsis Uzmanı

Firmanın bilanço ve gelir tablosu kalemlerinden 18 rasyoyu hesaplar, sizin verdiğiniz sektör değerleriyle karşılaştırır.
  Likidite   : cari oran, asit-test oranı, nakit oranı
  Mali yapı  : toplam borç / aktif, özkaynak / aktif, borç / özkaynak, kısa vadeli borç / toplam borç
  Kârlılık   : brüt, faaliyet ve net kâr marjı; aktif ve özkaynak kârlılığı
  Faaliyet   : alacak tahsil, stok tutma ve ticari borç ödeme süresi; nakit dönüş süresi; aktif devir hızı
  Borç servisi: faiz karşılama oranı = faaliyet kârı ÷ finansman giderleri
  Sapma %    = (firma − sektör) ÷ |sektör| × 100;  ±tolerans içindeyse "Sektöre yakın"
Sektör değerleri koda gömülü değildir (ör. TCMB Sektör Bilançoları'ndan siz girersiniz). Süreler dönem sonu bakiyesiyle hesaplanır.
Yorum cümleleri hazır kalıplardır; kredi kararı veya derecelendirme değildir.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


KALEM = {"donen": ("donen varliklar", "donen varliklar toplami", "toplam donen varliklar"), "hazir": ("hazir degerler", "nakit ve nakit benzerleri", "nakit ve benzerleri"),
         "alacak": ("ticari alacaklar",), "stok": ("stoklar",), "aktif": ("toplam varliklar", "aktif toplami", "toplam aktif", "varliklar toplami"),
         "kvyk": ("kisa vadeli yukumlulukler", "kisa vadeli yabanci kaynaklar", "kvyk"), "uvyk": ("uzun vadeli yukumlulukler", "uzun vadeli yabanci kaynaklar", "uvyk"),
         "ticari_borc": ("ticari borclar",), "ozkaynak": ("ozkaynaklar", "oz kaynaklar", "ozkaynak toplami", "toplam ozkaynaklar"),
         "satis": ("net satislar", "hasilat", "satis gelirleri"), "smm": ("satislarin maliyeti", "smm"),
         "faaliyet": ("faaliyet kari", "esas faaliyet kari", "faaliyet kari zarari", "esas faaliyet kari zarari"), "finansman": ("finansman giderleri",),
         "net": ("donem net kari", "net kar", "net donem kari", "donem net kari zarari", "net donem kari zarari")}
KALEM_AD = {"donen": "Dönen Varlıklar", "hazir": "Hazır Değerler", "alacak": "Ticari Alacaklar", "stok": "Stoklar", "aktif": "Toplam Varlıklar", "kvyk": "Kısa Vadeli Yükümlülükler",
            "uvyk": "Uzun Vadeli Yükümlülükler", "ticari_borc": "Ticari Borçlar", "ozkaynak": "Özkaynaklar", "satis": "Net Satışlar", "smm": "Satışların Maliyeti",
            "faaliyet": "Faaliyet Kârı", "finansman": "Finansman Giderleri", "net": "Dönem Net Kârı"}
G = 365


def bol(a, b, carpan=1):
    return None if a is None or b is None or b == 0 else yuvarla(Decimal(a) / Decimal(b) * carpan)


def _borc(k):
    return None if k["kvyk"] is None or k["uvyk"] is None else k["kvyk"] + k["uvyk"]


def _poz(x):
    return x if x is not None and x > 0 else None          # özkaynak negatifse oran anlamsızdır


def _topla(*x):
    return None if any(v is None for v in x) else sum(x)


# (kod, ad, grup, birim, yön, hesap, olumsuz yorum, olumlu yorum)   yön: yuksek = yüksek olması olumlu, dusuk = düşük olması olumlu, notr = yön verilmez
RASYO = [
    ("cari", "Cari Oran", "Likidite", "x", "yuksek", lambda k: bol(k["donen"], k["kvyk"]),
     "Kısa vadeli borçları dönen varlıklarla karşılama gücü sektöre göre zayıf.", "Kısa vadeli borç ödeme gücü sektörün üzerinde."),
    ("asit", "Asit-Test Oranı", "Likidite", "x", "yuksek", lambda k: bol(_topla(k["donen"], -k["stok"] if k["stok"] is not None else None), k["kvyk"]),
     "Stoklar hariç likidite sektöre göre zayıf; borç ödeme gücü stokların satışına bağlı.", "Stoklar hariç likidite sektörün üzerinde."),
    ("nakit", "Nakit Oranı", "Likidite", "x", "yuksek", lambda k: bol(k["hazir"], k["kvyk"]),
     "Hazır değerlerin kısa vadeli borçlara oranı sektörün altında.", "Hazır değerlerin kısa vadeli borçlara oranı sektörün üzerinde."),
    ("kaldirac", "Toplam Borç / Toplam Aktif (%)", "Mali Yapı", "%", "dusuk", lambda k: bol(_borc(k), k["aktif"], 100),
     "Varlıkların borçla finanse edilen kısmı sektörden yüksek.", "Borçluluk sektörün altında."),
    ("oz_aktif", "Özkaynak / Toplam Aktif (%)", "Mali Yapı", "%", "yuksek", lambda k: bol(k["ozkaynak"], k["aktif"], 100),
     "Özkaynakların aktif içindeki payı sektörün altında.", "Özkaynakların aktif içindeki payı sektörün üzerinde."),
    ("borc_oz", "Borç / Özkaynak", "Mali Yapı", "x", "dusuk", lambda k: bol(_borc(k), _poz(k["ozkaynak"])),
     "Özkaynağa göre borç yükü sektörden yüksek.", "Özkaynağa göre borç yükü sektörün altında."),
    ("kv_borc", "Kısa Vadeli Borç / Toplam Borç (%)", "Mali Yapı", "%", "dusuk", lambda k: bol(k["kvyk"], _borc(k), 100),
     "Borçların vadesi sektöre göre daha kısa; yenileme (çevirme) riski daha yüksek.", "Borçların vade yapısı sektöre göre daha uzun."),
    ("brut", "Brüt Kâr Marjı (%)", "Kârlılık", "%", "yuksek", lambda k: bol(_topla(k["satis"], -k["smm"] if k["smm"] is not None else None), k["satis"], 100),
     "Brüt kârlılık sektörün altında; fiyatlama gücü veya üretim maliyeti incelenmeli.", "Brüt kârlılık sektörün üzerinde."),
    ("faaliyet", "Faaliyet Kâr Marjı (%)", "Kârlılık", "%", "yuksek", lambda k: bol(k["faaliyet"], k["satis"], 100),
     "Faaliyet kârlılığı sektörün altında.", "Faaliyet kârlılığı sektörün üzerinde."),
    ("net", "Net Kâr Marjı (%)", "Kârlılık", "%", "yuksek", lambda k: bol(k["net"], k["satis"], 100),
     "Net kârlılık sektörün altında.", "Net kârlılık sektörün üzerinde."),
    ("roa", "Aktif Kârlılığı (%)", "Kârlılık", "%", "yuksek", lambda k: bol(k["net"], k["aktif"], 100),
     "Varlıkların kâr üretme gücü sektörün altında.", "Varlıkların kâr üretme gücü sektörün üzerinde."),
    ("roe", "Özkaynak Kârlılığı (%)", "Kârlılık", "%", "yuksek", lambda k: bol(k["net"], _poz(k["ozkaynak"]), 100),
     "Özkaynak kârlılığı sektörün altında.", "Özkaynak kârlılığı sektörün üzerinde."),
    ("alacak_sure", "Alacak Tahsil Süresi (gün)", "Faaliyet", "gün", "dusuk", lambda k: bol(k["alacak"], k["satis"], G),
     "Alacaklar sektöre göre daha geç tahsil ediliyor.", "Alacaklar sektöre göre daha hızlı tahsil ediliyor."),
    ("stok_sure", "Stok Tutma Süresi (gün)", "Faaliyet", "gün", "dusuk", lambda k: bol(k["stok"], k["smm"], G),
     "Stoklar sektöre göre daha uzun süre elde tutuluyor.", "Stoklar sektöre göre daha hızlı dönüyor."),
    ("borc_sure", "Ticari Borç Ödeme Süresi (gün)", "Faaliyet", "gün", "notr", lambda k: bol(k["ticari_borc"], k["smm"], G),
     "", ""),
    ("nakit_donus", "Nakit Dönüş Süresi (gün)", "Faaliyet", "gün", "dusuk",
     lambda k: _topla(bol(k["alacak"], k["satis"], G), bol(k["stok"], k["smm"], G), -bol(k["ticari_borc"], k["smm"], G) if bol(k["ticari_borc"], k["smm"], G) is not None else None),
     "Nakit dönüş süresi sektörden uzun; işletme sermayesi ihtiyacı daha yüksek.", "Nakit dönüş süresi sektörden kısa."),
    ("aktif_devir", "Aktif Devir Hızı", "Faaliyet", "x", "yuksek", lambda k: bol(k["satis"], k["aktif"]),
     "Varlıkların satış üretme hızı sektörün altında.", "Varlıkların satış üretme hızı sektörün üzerinde."),
    ("faiz", "Faiz Karşılama Oranı", "Borç Servisi", "x", "yuksek", lambda k: bol(k["faaliyet"], k["finansman"]),
     "Faaliyet kârının finansman giderini karşılama gücü sektörün altında.", "Faaliyet kârının finansman giderini karşılama gücü sektörün üzerinde."),
]
RASYO_ES = {"asit": ("asit test orani", "likidite orani"), "kaldirac": ("toplam borc toplam aktif", "kaldirac orani", "yabanci kaynaklar aktif toplami"),
            "oz_aktif": ("ozkaynak toplam aktif", "ozkaynaklar aktif toplami", "ozkaynak orani"), "borc_oz": ("borc ozkaynak", "yabanci kaynaklar ozkaynaklar"),
            "kv_borc": ("kisa vadeli borc toplam borc",), "roa": ("aktif karliligi", "roa"), "roe": ("ozkaynak karliligi", "roe"),
            "alacak_sure": ("alacak tahsil suresi", "ortalama tahsil suresi"), "stok_sure": ("stok tutma suresi", "stok devir suresi"),
            "borc_sure": ("ticari borc odeme suresi", "borc odeme suresi"), "faiz": ("faiz karsilama orani", "faiz karsilama")}


def _ad_katla(s) -> str:
    return katla(re.sub(r"\((%|gün|gun|x)\)", "", metin(s)))


def firma_oku(yol: Path) -> tuple[list[str], list[dict], list[dict]]:
    """Kalem | dönem 1 | dönem 2 ...  →  (dönem adları, dönem başına kalem sözlüğü, uyarılar)"""
    s, uy = tablo_oku(yol), []
    bi = next((i for i, r in enumerate(s[:15]) if katla(r[0]) in ("kalem", "hesap", "hesap adi", "kalem adi") and len(r) > 1), None)
    if bi is None:
        raise ValueError(f"{yol.name}: başlık satırı bulunamadı; ilk sütun 'Kalem', sonraki sütunlar dönem adları olmalı.")
    donemler = [(j, metin(c)) for j, c in enumerate(s[bi][1:], 1) if metin(c)]
    if not donemler:
        raise ValueError(f"{yol.name}: dönem sütunu yok.")
    ters = {a: kod for kod, adlar in KALEM.items() for a in adlar}
    veri = [{k: None for k in KALEM} for _ in donemler]
    for n, r in enumerate(s[bi + 1:], bi + 2):
        kod = ters.get(katla(r[0]))
        if not kod:
            continue
        for i, (j, ad) in enumerate(donemler):
            h = r[j] if j < len(r) else None
            v = para(h)
            if v is None and metin(h):
                uy.append(uyari("Yüksek", "Tutar okunamadı", f"{KALEM_AD[kod]} · {ad}", f"{yol.name} satır {n}: '{metin(h)}'. Bu kaleme bağlı rasyolar hesaplanmadı."))
            if v is not None and kod in ("smm", "finansman"):
                v = abs(v)                                    # gider kalemleri eksi işaretli girilmiş olabilir
            veri[i][kod] = v
    for i, (_, ad) in enumerate(donemler):
        k = veri[i]
        eksik = [KALEM_AD[x] for x in KALEM if k[x] is None]
        if eksik:
            uy.append(uyari("Orta", "Eksik kalem", ad, f"{', '.join(eksik)} bulunamadı; bu kalemlere bağlı rasyolar boş kalır."))
        if None not in (k["aktif"], k["kvyk"], k["uvyk"], k["ozkaynak"]) and abs(k["aktif"] - k["kvyk"] - k["uvyk"] - k["ozkaynak"]) > 1:
            uy.append(uyari("Yüksek", "Bilanço denk değil", ad, f"Toplam varlıklar {k['aktif']:f} ≠ yükümlülükler + özkaynaklar {k['kvyk'] + k['uvyk'] + k['ozkaynak']:f}."))
    return [d[1] for d in donemler], veri, uy


def sektor_oku(yol: Path) -> tuple[dict, list[dict]]:
    uy, sektor = [], {}
    es = {_ad_katla(r[1]): r[0] for r in RASYO} | {a: kod for kod, adlar in RASYO_ES.items() for a in adlar}
    for r in kayitlar(yol, {"rasyo": ("rasyo", "oran", "gosterge"), "deger": ("sektor ortalamasi", "sektor", "sektor degeri", "ortalama", "medyan"),
                            "q1": ("alt ceyrek", "q1", "1 ceyrek"), "q3": ("ust ceyrek", "q3", "3 ceyrek")}, ("rasyo", "deger")):
        if not metin(r["rasyo"]):
            continue
        kod, v = es.get(_ad_katla(r["rasyo"])), para(r["deger"])
        if not kod:
            uy.append(uyari("Orta", "Sektör rasyosu tanınmadı", metin(r["rasyo"]), f"{yol.name} satır {r['_satir']}: bu ad paketteki rasyolarla eşleşmedi; README'deki adları kullanın."))
        elif v is None:
            uy.append(uyari("Orta", "Sektör değeri okunamadı", metin(r["rasyo"]), f"{yol.name} satır {r['_satir']}: '{metin(r['deger'])}'."))
        else:
            q1, q3 = para(r["q1"]), para(r["q3"])
            if q1 is not None and q3 is not None and q1 > q3:
                uy.append(uyari("Orta", "Çeyrek değerleri ters", metin(r["rasyo"]), f"Alt çeyrek {q1:f} > üst çeyrek {q3:f}; çeyrek konumu hesaplanmadı."))
                q1 = q3 = None
            sektor[kod] = {"deger": v, "q1": q1, "q3": q3}
    return sektor, uy


def hesapla(k: dict) -> dict:
    return {r[0]: r[5](k) for r in RASYO}


def karsilastir(donemler: list[str], veri: list[dict], sektor: dict, tolerans: Decimal = Decimal(10)) -> dict:
    rasyolar = [hesapla(k) for k in veri]
    son, satirlar, notlar = rasyolar[-1], [], []
    for kod, ad, grup, birim, yon, _, olumsuz, olumlu in RASYO:
        f, s = son[kod], sektor.get(kod)
        x = {"kod": kod, "ad": ad, "grup": grup, "birim": birim, "firma": f, "sektor": s["deger"] if s else None, "sapma": None, "durum": "", "ceyrek": "", "yorum": "", "seyir": ""}
        if len(rasyolar) > 1 and f is not None and rasyolar[-2][kod] is not None:
            o = rasyolar[-2][kod]
            x["seyir"] = "Değişmedi" if f == o else "Arttı" if yon == "notr" and f > o else "Azaldı" if yon == "notr" else "İyileşti" if (f > o) == (yon == "yuksek") else "Kötüleşti"
        if f is None:
            x["durum"] = "Hesaplanamadı"
        elif not s:
            x["durum"] = "Sektör verisi yok"
        else:
            sv = s["deger"]
            x["sapma"] = yuvarla((f - sv) / abs(sv) * 100) if sv else None
            fark = "üstünde" if f > sv else "altında"
            if f == sv or (x["sapma"] is not None and abs(x["sapma"]) <= tolerans):
                x["durum"] = "Sektöre yakın"
            elif yon == "notr":
                x["durum"] = "Sektörden farklı"
                x["yorum"] = f"{ad}: firma {f}, sektör {sv} (sektörün {fark}). Yönü tek başına olumlu veya olumsuz değildir; tedarikçi vadeleriyle birlikte değerlendirin."
            else:
                iyi = (f > sv) == (yon == "yuksek")
                x["durum"] = "Olumlu sapma" if iyi else "Olumsuz sapma"
                sapma = f"%{abs(x['sapma'])} " if x["sapma"] is not None else ""
                x["yorum"] = f"{ad}: firma {f}, sektör {sv} (sektörün {sapma}{fark}). {olumlu if iyi else olumsuz}"
            if s["q1"] is not None and s["q3"] is not None:
                x["ceyrek"] = "Alt çeyreğin altında" if f < s["q1"] else "Üst çeyreğin üstünde" if f > s["q3"] else "Çeyrekler arasında"
        satirlar.append(x)
    k = veri[-1]
    if k["ozkaynak"] is not None and k["ozkaynak"] <= 0:
        notlar.append(("Yüksek", "Özkaynak sıfır veya negatif", "Borç / özkaynak ve özkaynak kârlılığı hesaplanmadı; firma teknik olarak borca batık görünüyor olabilir."))
    if son["faiz"] is not None and son["faiz"] < 1:
        notlar.append(("Yüksek", "Faaliyet kârı finansman giderini karşılamıyor", f"Faiz karşılama oranı {son['faiz']} (1'in altında)."))
    if k["net"] is not None and k["net"] < 0:
        notlar.append(("Orta", "Dönem zararı", "Son dönemde net zarar var; kârlılık rasyoları negatif."))
    if son["cari"] is not None and son["cari"] < 1:
        notlar.append(("Orta", "Net işletme sermayesi negatif", f"Cari oran {son['cari']} (1'in altında): kısa vadeli yükümlülükler dönen varlıklardan büyük."))
    return {"donemler": donemler, "rasyolar": rasyolar, "satirlar": satirlar, "notlar": notlar, "tolerans": tolerans}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Karşılaştırma"
    ws.append([f"Dönem: {r['donemler'][-1]} · Tolerans ±%{r['tolerans']}"])
    ws["A1"].font = Font(bold=True)
    _baslik(ws, ["Grup", "Rasyo", "Firma", "Sektör", "Sapma %", "Durum", "Çeyrek Konumu", "Önceki Döneme Göre"], (14, 36, 12, 12, 11, 20, 22, 18), "A3")
    renk = {"Olumsuz sapma": KIRMIZI, "Olumlu sapma": YESIL, "Sektörden farklı": SARI}
    for x in r["satirlar"]:
        _satir(ws, [x["grup"], x["ad"], x["firma"], x["sektor"], x["sapma"], x["durum"], x["ceyrek"], x["seyir"]], para_s=(3, 4, 5),
               dolgu={6: renk.get(x["durum"]), 8: KIRMIZI if x["seyir"] == "Kötüleşti" else None})
    ws = wb.create_sheet("Dönemsel Rasyolar")
    _baslik(ws, ["Grup", "Rasyo", *r["donemler"]], (14, 36, *[14] * len(r["donemler"])))
    for i, (kod, ad, grup, *_) in enumerate(RASYO):
        _satir(ws, [grup, ad, *[d[kod] for d in r["rasyolar"]]], para_s=range(3, 3 + len(r["donemler"])))
    ws = wb.create_sheet("Yorumlar")
    _baslik(ws, ["Tür", "Konu", "Açıklama", "Analist Notu (siz doldurun)"], (16, 40, 100, 40))
    for d, konu, ac in r["notlar"]:
        _satir(ws, ["Dikkat", konu, ac, ""], dolgu={1: RENK[d]}, sar=(3,))
    for durum in ("Olumsuz sapma", "Sektörden farklı", "Olumlu sapma"):
        for x in r["satirlar"]:
            if x["durum"] == durum:
                _satir(ws, [durum, x["ad"], x["yorum"], ""], dolgu={1: renk[durum]}, sar=(3,))
    ws.append([])
    ws.append(["", "", "Yorumlar hazır kalıplardır; rakamın nedenini açıklamaz. Sektör değerinin kaynağını, dönemini ve tanımını ayrıca not edin."])
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(firma: Path, sektor_yol: Path, cikti: Path, tolerans: Decimal = Decimal(10)) -> dict:
    donemler, veri, uy = firma_oku(firma)
    sektor, uy2 = sektor_oku(sektor_yol)
    r = karsilastir(donemler, veri, sektor, tolerans)
    r["uyarilar"] = sirala(uy + uy2)
    rapor_yaz(cikti, r, uy + uy2)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Firma rasyolarını hesaplayıp verilen sektör değerleriyle karşılaştırır.")
    p.add_argument("--firma", type=Path, default=ORNEK / "firma_finansallari.csv", help="İlk sütun Kalem, sonraki sütunlar dönemler (son sütun karşılaştırılır)")
    p.add_argument("--sektor", type=Path, default=ORNEK / "sektor_degerleri.csv", help="Rasyo, Sektör Ortalaması, Alt Çeyrek, Üst Çeyrek")
    p.add_argument("--tolerans", default="10", help="Sektöre yakın sayılacak sapma %% (varsayılan 10)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "rasyo_karsilastirma.xlsx")
    a = p.parse_args(argv)
    for y in (a.firma, a.sektor):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    tol = para(a.tolerans)
    if tol is None or tol < 0:
        print("[X] --tolerans sıfır veya pozitif bir sayı olmalı")
        return 2
    try:
        r = calistir(a.firma, a.sektor, a.cikti, tol)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    say = lambda d: sum(x["durum"] == d for x in r["satirlar"])   # noqa: E731
    print(f"[OK] {r['donemler'][-1]} · {say('Olumsuz sapma')} olumsuz, {say('Olumlu sapma')} olumlu, {say('Sektöre yakın')} sektöre yakın, {say('Sektör verisi yok')} rasyoda sektör verisi yok")
    for d, konu, ac in r["notlar"]:
        print(f"[!] {konu}: {ac}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
