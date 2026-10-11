import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def kalem(**kw):
    k = dict(donen=1800, hazir=80, alacak=900, stok=720, aktif=3000, kvyk=1500, uvyk=600, ticari_borc=600, ozkaynak=900, satis=4500, smm=3600, faaliyet=270, finansman=300, net=-45)
    return {a: (D(v) if v is not None else None) for a, v in (k | kw).items()}


def sek(**kw):
    return {k: {"deger": D(str(v[0])), "q1": D(str(v[1])) if len(v) > 1 else None, "q3": D(str(v[2])) if len(v) > 1 else None} for k, v in kw.items()}


class RasyoTesti(unittest.TestCase):
    def test_rasyolar(self):
        r = main.hesapla(kalem())
        beklenen = dict(cari="1.20", asit="0.72", nakit="0.05", kaldirac="70.00", oz_aktif="30.00", borc_oz="2.33", kv_borc="71.43", brut="20.00", faaliyet="6.00", net="-1.00",
                        roa="-1.50", roe="-5.00", alacak_sure="73.00", stok_sure="73.00", borc_sure="60.83", nakit_donus="85.17", aktif_devir="1.50", faiz="0.90")
        self.assertEqual(r, {k: D(v) for k, v in beklenen.items()})
        self.assertEqual(len(main.RASYO), 18)

    def test_eksik_ve_negatif_ozkaynak(self):
        r = main.hesapla(kalem(stok=None, ozkaynak=-100, finansman=0))
        self.assertEqual((r["asit"], r["stok_sure"], r["nakit_donus"], r["borc_oz"], r["roe"], r["faiz"], r["cari"]), (None, None, None, None, None, None, D("1.20")))
        self.assertEqual(r["oz_aktif"], D("-3.33"))

    def test_karsilastirma(self):
        s = sek(cari=(1.45, 1.05, 1.90), kaldirac=(62, 48, 74), brut=(21,), net=(3.2, 0.5, 6), borc_sure=(40,), aktif_devir=(1.2, 0.8, 1.4), stok_sure=(73,))
        r = main.karsilastir(["2025"], [kalem()], s)
        d = {x["kod"]: (x["durum"], x["sapma"], x["ceyrek"]) for x in r["satirlar"]}
        self.assertEqual(d["cari"], ("Olumsuz sapma", D("-17.24"), "Çeyrekler arasında"))
        self.assertEqual(d["kaldirac"], ("Olumsuz sapma", D("12.90"), "Çeyrekler arasında"))      # düşük olması olumlu
        self.assertEqual(d["brut"], ("Sektöre yakın", D("-4.76"), ""))
        self.assertEqual(d["net"], ("Olumsuz sapma", D("-131.25"), "Alt çeyreğin altında"))
        self.assertEqual(d["borc_sure"], ("Sektörden farklı", D("52.08"), ""))
        self.assertEqual(d["aktif_devir"], ("Olumlu sapma", D("25.00"), "Üst çeyreğin üstünde"))
        self.assertEqual(d["stok_sure"], ("Sektöre yakın", D("0.00"), ""))
        self.assertEqual(d["roe"], ("Sektör verisi yok", None, ""))
        y = {x["kod"]: x["yorum"] for x in r["satirlar"]}
        self.assertIn("firma 1.20, sektör 1.45 (sektörün %17.24 altında)", y["cari"])
        self.assertEqual((y["brut"], y["roe"]), ("", ""))
        self.assertEqual(main.karsilastir(["2025"], [kalem()], s, D(20))["satirlar"][0]["durum"], "Sektöre yakın")
        self.assertEqual(main.karsilastir(["2025"], [kalem()], sek(cari=(1.3334,)), D(10))["satirlar"][0]["sapma"], D("-10.00"))   # sınır dahil

    def test_notlar_ve_seyir(self):
        r = main.karsilastir(["2024", "2025"], [kalem(donen=1300, kvyk=900, uvyk=500, ozkaynak=1000, aktif=2400, net=100, faaliyet=320, finansman=160, stok=500, satis=4000, smm=3200,
                                                         alacak=600, ticari_borc=400), kalem(kvyk=2000, uvyk=100)], {})
        self.assertEqual([n[1] for n in r["notlar"]], ["Faaliyet kârı finansman giderini karşılamıyor", "Dönem zararı", "Net işletme sermayesi negatif"])
        d = {x["kod"]: x["seyir"] for x in r["satirlar"]}
        self.assertEqual((d["cari"], d["kaldirac"], d["borc_sure"], d["brut"], d["faiz"]), ("Kötüleşti", "Kötüleşti", "Arttı", "Değişmedi", "Kötüleşti"))
        self.assertEqual([n[1] for n in main.karsilastir(["x"], [kalem(ozkaynak=0, net=5, faaliyet=400)], {})["notlar"]], ["Özkaynak sıfır veya negatif"])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "r.xlsx"
            r = main.calistir(ORNEK / "firma_finansallari.csv", ORNEK / "sektor_degerleri.csv", cikti)
            self.assertEqual(r["rasyolar"][0]["cari"], D("1.44"))
            self.assertEqual((r["rasyolar"][0]["roe"], r["rasyolar"][0]["stok_sure"], r["rasyolar"][0]["faiz"]), (D("10.00"), D("57.03"), D("2.00")))
            self.assertEqual(r["rasyolar"][1], main.hesapla(kalem()))                               # örnek 2025 = testteki kalemlerin 10.000 katı
            say = lambda d: sum(x["durum"] == d for x in r["satirlar"])   # noqa: E731
            self.assertEqual((say("Olumsuz sapma"), say("Sektöre yakın"), say("Sektör verisi yok"), say("Olumlu sapma")), (10, 2, 6, 0))
            self.assertEqual(r["uyarilar"], [])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Karşılaştırma", "Dönemsel Rasyolar", "Yorumlar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            f, s = Path(tmp) / "f.csv", Path(tmp) / "s.csv"
            f.write_text("Firma X\nKalem;2025\nDönen Varlıklar;1.000\nKVYK;500\nUVYK;200\nÖzkaynaklar;200\nAktif Toplamı;1.000\nStoklar;çok\nSatışların Maliyeti (-);-400\n", encoding="utf-8")
            s.write_text("Rasyo;Sektör;Alt Çeyrek;Üst Çeyrek\nCari Oran;1,5\nLikidite Oranı;0,9\nROE;12\nBilinmeyen;1\nNakit Oranı;yok\nNet Kâr Marjı (%);3;5;1\n", encoding="utf-8")
            d, v, uy = main.firma_oku(f)
            self.assertEqual((d, v[0]["donen"], v[0]["smm"], v[0]["stok"]), (["2025"], D(1000), D(400), None))
            self.assertEqual(sorted(u["tur"] for u in uy), ["Bilanço denk değil", "Eksik kalem", "Tutar okunamadı"])
            sk, uy = main.sektor_oku(s)
            self.assertEqual((sorted(sk), sk["net"]["q1"]), (["asit", "cari", "net", "roe"], None))
            self.assertEqual(sorted(u["tur"] for u in uy), ["Sektör değeri okunamadı", "Sektör rasyosu tanınmadı", "Çeyrek değerleri ters"])
            c = str(Path(tmp) / "x.xlsx")
            self.assertEqual(main.main(["--firma", str(f), "--sektor", str(s), "--cikti", c]), 0)
            self.assertEqual(main.main(["--firma", str(f), "--sektor", str(s), "--cikti", c, "--tolerans", "az"]), 2)
            self.assertEqual(main.main(["--firma", str(Path(tmp) / "yok.csv")]), 1)
            f.write_text("Hesap Kodu;Tutar\n1;2\n", encoding="utf-8")
            self.assertEqual(main.main(["--firma", str(f), "--cikti", c]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
