"""
İş Kazası ve Ramak Kala Analizi — Gerisi Bende kod bloğu
İş Sağlığı ve Güvenliği › İş Güvenliği Uzmanı

Kaza ve ramak kala kayıtlarından bölüm, neden ve zaman bazlı analiz ile kaza sıklık ve ağırlık oranlarını üretir.
  Kaza sıklık oranı          = iş kazası sayısı × 1.000.000 ÷ toplam çalışma saati
  Kaza ağırlık oranı         = kayıp iş günü × 1.000 ÷ toplam çalışma saati
  1.000 çalışan başına kaza  = iş kazası sayısı × 1.000 ÷ ortalama çalışan sayısı
  Ramak kala / kaza          = ramak kala sayısı ÷ iş kazası sayısı
  Kaza başına kayıp gün      = kayıp iş günü ÷ iş kazası sayısı
  Neden Pareto               = nedenler olay sayısına göre azalan sırada; pay % ve kümülatif %
Çarpanlar (1.000.000 ve 1.000) komut satırından değiştirilebilir; kıyasladığınız kaynak hangi çarpanı kullanıyorsa onu verin.
Oranlar yalnız "İş Kazası" kayıtlarından hesaplanır; ramak kala olayları orana girmez.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
from datetime import time

O_SUTUN = {"tarih": ("tarih", "olay tarihi", "kaza tarihi"), "saat": ("saat", "olay saati"), "tur": ("tur", "olay turu"), "bolum": ("bolum", "birim", "departman"), "neden": ("neden", "kok neden", "olay nedeni", "kaza nedeni"),
           "kayip": ("kayip gun", "kayip is gunu", "is goremezlik gunu", "rapor gunu"), "vardiya": ("vardiya",), "aciklama": ("aciklama", "olay ozeti")}
B_SUTUN = {"bolum": ("bolum", "birim", "departman"), "calisan": ("calisan sayisi", "calisan", "kisi sayisi", "personel"), "saat": ("calisma saati", "toplam calisma saati", "saat")}
TUR = {"is kazasi": "İş Kazası", "kaza": "İş Kazası", "ramak kala": "Ramak Kala", "ramakkala": "Ramak Kala", "ramak kala olay": "Ramak Kala"}
GUNLER = ("Pazartesi", "Salı", "Çarşamba", "Perşembe", "Cuma", "Cumartesi", "Pazar")
SAAT_DILIM = ("00–04", "04–08", "08–12", "12–16", "16–20", "20–24")
Z = Decimal(0)


def saat_oku(x) -> int | None:
    if isinstance(x, (datetime, time)):
        return x.hour
    m = re.fullmatch(r"\s*(\d{1,2})(?:[:.](\d{2}))?(?::\d{2})?\s*", str(x or ""))
    return int(m.group(1)) if m and int(m.group(1)) < 24 and int(m.group(2) or 0) < 60 else None


def oku_olaylar(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, O_SUTUN, ("tarih", "tur")):
        if not metin(r["tarih"]) and not metin(r["tur"]):
            continue
        t, tur, kayip = tarih(r["tarih"]), TUR.get(katla(r["tur"])), para(r["kayip"])
        if t is None or tur is None or (metin(r["kayip"]) and (kayip is None or kayip < 0)):
            uy.append(uyari("Yüksek", "Olay okunamadı", f"Satır {r['_satir']}", f"Tarih '{metin(r['tarih'])}', tür '{metin(r['tur'])}', kayıp gün '{metin(r['kayip'])}'. Tür 'İş Kazası' ya da 'Ramak Kala' olmalı. Satır alınmadı."))
            continue
        l.append({"tarih": t, "saat": saat_oku(r["saat"]), "tur": tur, "bolum": metin(r["bolum"]) or "(bölüm yok)", "neden": metin(r["neden"]) or "(neden yok)", "kayip": kayip, "vardiya": metin(r["vardiya"]),
                  "aciklama": metin(r["aciklama"]), "satir": r["_satir"]})
    return l, uy


def oku_bolumler(yol: Path) -> list[dict]:
    l = []
    for r in kayitlar(yol, B_SUTUN, ("bolum",)):
        if not metin(r["bolum"]):
            continue
        c, s = para(r["calisan"]), para(r["saat"])
        if (c is not None and c < 0) or (s is not None and s < 0) or (metin(r["calisan"]) and c is None) or (metin(r["saat"]) and s is None):
            raise ValueError(f"{yol.name} satır {r['_satir']}: çalışan sayısı ve çalışma saati sıfır veya pozitif sayı olmalı.")
        l.append({"bolum": metin(r["bolum"]), "calisan": c, "saat": s})
    return l


def _oran(pay, payda, carpan=1) -> Decimal | None:
    return yuvarla(Decimal(pay) * carpan / Decimal(payda)) if payda else None


def analiz(olaylar: list[dict], saat: Decimal | None = None, calisan: Decimal | None = None, bolum_bilgi: list[dict] | None = None, siklik_carpan: Decimal = Decimal(1000000), agirlik_carpan: Decimal = Decimal(1000),
           tekrar: int = 3) -> dict:
    if not olaylar:
        raise ValueError("Analiz edilecek olay yok.")
    uy = []
    bilgi = {katla(b["bolum"]): b for b in bolum_bilgi or []}
    if bilgi:
        saat = saat if saat is not None else sum((b["saat"] or Z for b in bilgi.values()), Z) or None
        calisan = calisan if calisan is not None else sum((b["calisan"] or Z for b in bilgi.values()), Z) or None

    def say(liste: list[dict]) -> dict:
        kazalar = [o for o in liste if o["tur"] == "İş Kazası"]
        return {"kaza": len(kazalar), "ramak": len(liste) - len(kazalar), "kayip": sum((o["kayip"] or Z for o in kazalar), Z)}

    def kirilim(anahtar, sira=None) -> list[dict]:
        grup = {}
        for o in olaylar:
            ad = anahtar(o)
            if ad is not None:
                grup.setdefault(katla(ad), (ad, []))[1].append(o)
        satirlar = [{"ad": ad, **say(liste)} for ad, liste in grup.values()]
        if sira:
            var = {s["ad"]: s for s in satirlar}
            return [var.get(ad, {"ad": ad, "kaza": 0, "ramak": 0, "kayip": Z}) for ad in sira] + [s for s in satirlar if s["ad"] not in sira]
        return sorted(satirlar, key=lambda s: katla(s["ad"]))

    t = say(olaylar)
    kazalar = [o for o in olaylar if o["tur"] == "İş Kazası"]
    ozet = {**t, "olay": len(olaylar), "kayipli_kaza": sum(1 for o in kazalar if o["kayip"]), "saat": saat, "calisan": calisan, "siklik": _oran(t["kaza"], saat, siklik_carpan), "agirlik": _oran(t["kayip"], saat, agirlik_carpan),
            "bin_calisan": _oran(t["kaza"], calisan, 1000), "ramak_kaza": _oran(t["ramak"], t["kaza"]), "ort_kayip": _oran(t["kayip"], t["kaza"])}
    bolumler = kirilim(lambda o: o["bolum"])
    for ad in (b["bolum"] for b in bilgi.values()):
        if katla(ad) not in {katla(b["ad"]) for b in bolumler}:
            bolumler.append({"ad": ad, "kaza": 0, "ramak": 0, "kayip": Z})
    bolumler.sort(key=lambda b: (-b["kaza"], -b["ramak"], katla(b["ad"])))
    for b in bolumler:
        bb = bilgi.get(katla(b["ad"]), {})
        b["calisan"], b["saat"] = bb.get("calisan"), bb.get("saat")
        b["siklik"], b["agirlik"] = _oran(b["kaza"], b["saat"], siklik_carpan), _oran(b["kayip"], b["saat"], agirlik_carpan)
        if bilgi and not bb:
            uy.append(uyari("Orta", "Bölüm bilgisi yok", b["ad"], f"Olay kaydında geçiyor ({b['kaza']} kaza, {b['ramak']} ramak kala) ama bölüm dosyasında yok; bölüm oranı hesaplanamadı ve toplam çalışma saatine katılmadı."))
    nedenler = sorted(kirilim(lambda o: o["neden"]), key=lambda n: (-(n["kaza"] + n["ramak"]), -n["kaza"], katla(n["ad"])))
    kum = 0
    for n in nedenler:
        n["toplam"] = n["kaza"] + n["ramak"]
        kum += n["toplam"]
        n["pay"], n["kumulatif"] = _oran(n["toplam"] * 100, len(olaylar)), _oran(kum * 100, len(olaylar))
        if n["toplam"] >= tekrar and n["ad"] != "(neden yok)":
            uy.append(uyari("Orta", "Tekrarlayan neden", n["ad"], f"Bu nedenle {n['toplam']} olay kayıtlı ({n['kaza']} kaza, {n['ramak']} ramak kala; kayıp gün {n['kayip']:f}). Düzeltici faaliyetin etkisini gözden geçirin."))
    if t["kaza"] and not t["ramak"]:
        uy.append(uyari("Bilgi", "Ramak kala kaydı yok", f"{t['kaza']} kaza", "Dönemde kaza var ama hiç ramak kala kaydı yok. Ramak kala olayları bildirilmiyor olabilir."))
    for o in kazalar:
        if o["kayip"] is None:
            uy.append(uyari("Bilgi", "Kayıp gün yazılmamış", f"{o['tarih']:%d.%m.%Y} · {o['bolum']}", f"Satır {o['satir']}: iş kazasının kayıp gün sayısı boş; sıfır sayıldı. İş göremezlik yoksa 0 yazın."))
    if saat is None:
        uy.append(uyari("Orta", "Çalışma saati verilmedi", "Oranlar", "Toplam çalışma saati verilmediği için kaza sıklık ve ağırlık oranları hesaplanamadı (--saat ya da --bolumler)."))
    return {"olaylar": olaylar, "ozet": ozet, "bolumler": bolumler, "nedenler": nedenler, "aylar": kirilim(lambda o: f"{o['tarih']:%Y-%m}"), "gunler": kirilim(lambda o: GUNLER[o["tarih"].weekday()], GUNLER),
            "saatler": kirilim(lambda o: SAAT_DILIM[o["saat"] // 4] if o["saat"] is not None else "(saat yok)", SAAT_DILIM), "vardiyalar": kirilim(lambda o: o["vardiya"] or None), "carpanlar": (siklik_carpan, agirlik_carpan), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    o, (sc, ac) = r["ozet"], r["carpanlar"]
    _baslik(ws, ["Gösterge", "Değer"], (52, 16))
    for etiket, deger, para_mi in (("Toplam olay", o["olay"], False), ("İş kazası", o["kaza"], False), ("Kayıp günlü iş kazası", o["kayipli_kaza"], False), ("Ramak kala", o["ramak"], False), ("Kayıp iş günü", o["kayip"], True),
                                   ("Toplam çalışma saati", o["saat"], True), ("Ortalama çalışan sayısı", o["calisan"], True), (f"Kaza sıklık oranı (kaza × {sc:,.0f} ÷ saat)", o["siklik"], True),
                                   (f"Kaza ağırlık oranı (kayıp gün × {ac:,.0f} ÷ saat)", o["agirlik"], True), ("1.000 çalışan başına kaza", o["bin_calisan"], True), ("Ramak kala / kaza", o["ramak_kaza"], True),
                                   ("Kaza başına ortalama kayıp gün", o["ort_kayip"], True)):
        _satir(ws, [etiket, deger], para_s=(2,) if para_mi else ())
    ws = wb.create_sheet("Bölümler")
    _baslik(ws, ["Bölüm", "İş Kazası", "Ramak Kala", "Kayıp Gün", "Çalışan", "Çalışma Saati", "Sıklık Oranı", "Ağırlık Oranı"], (28, 10, 10, 10, 10, 15, 12, 12))
    for b in r["bolumler"]:
        _satir(ws, [b["ad"], b["kaza"], b["ramak"], b["kayip"], b["calisan"], b["saat"], b["siklik"], b["agirlik"]], para_s=(6, 7, 8))
    ws = wb.create_sheet("Neden Pareto")
    _baslik(ws, ["Neden", "Toplam Olay", "İş Kazası", "Ramak Kala", "Kayıp Gün", "Pay %", "Kümülatif %"], (36, 10, 10, 10, 10, 10, 11))
    for n in r["nedenler"]:
        _satir(ws, [n["ad"], n["toplam"], n["kaza"], n["ramak"], n["kayip"], n["pay"], n["kumulatif"]], para_s=(6, 7), dolgu={7: SARI if n["kumulatif"] - n["pay"] < 80 else None})
    for baslik, alan, ad in (("Aylar", "aylar", "Ay"), ("Haftanın Günü", "gunler", "Gün"), ("Saat Dilimi", "saatler", "Saat"), ("Vardiya", "vardiyalar", "Vardiya")):
        if r[alan]:
            ws = wb.create_sheet(baslik)
            _baslik(ws, [ad, "İş Kazası", "Ramak Kala", "Kayıp Gün"], (18, 10, 10, 10))
            for s in r[alan]:
                _satir(ws, [s["ad"], s["kaza"], s["ramak"], s["kayip"]])
    ws = wb.create_sheet("Olaylar")
    _baslik(ws, ["Tarih", "Gün", "Saat", "Tür", "Bölüm", "Neden", "Kayıp Gün", "Vardiya", "Açıklama"], (12, 12, 7, 12, 22, 30, 10, 12, 50))
    for x in sorted(r["olaylar"], key=lambda x: (x["tarih"], x["saat"] if x["saat"] is not None else -1)):
        _satir(ws, [x["tarih"], GUNLER[x["tarih"].weekday()], x["saat"], x["tur"], x["bolum"], x["neden"], x["kayip"], x["vardiya"], x["aciklama"]], tarih_s=(1,), dolgu={4: KIRMIZI if x["tur"] == "İş Kazası" else SARI}, sar=(9,))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(o_yol: Path, b_yol: Path | None, cikti: Path, **kw) -> dict:
    olaylar, uy = oku_olaylar(o_yol)
    r = analiz(olaylar, bolum_bilgi=oku_bolumler(b_yol) if b_yol else None, **kw)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Kaza ve ramak kala kayıtlarından bölüm, neden ve zaman bazlı analiz ile kaza sıklık ve ağırlık oranlarını üretir.")
    p.add_argument("--olaylar", type=Path, help="Tarih, Saat, Tür (İş Kazası / Ramak Kala), Bölüm, Neden, Kayıp Gün, Vardiya, Açıklama")
    p.add_argument("--bolumler", type=Path, help="Bölüm, Çalışan Sayısı, Çalışma Saati (isteğe bağlı; bölüm oranları için)")
    p.add_argument("--saat", help="Dönemdeki toplam çalışma saati (bölüm dosyası yoksa ya da toplamı ezmek için)")
    p.add_argument("--calisan", help="Dönemdeki ortalama çalışan sayısı")
    p.add_argument("--siklik-carpan", default="1000000", help="Sıklık oranı çarpanı (varsayılan 1.000.000)")
    p.add_argument("--agirlik-carpan", default="1000", help="Ağırlık oranı çarpanı (varsayılan 1.000)")
    p.add_argument("--tekrar", type=int, default=3, help="Aynı nedenle bu kadar ve üstü olay varsa uyarılır (varsayılan 3)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "is_kazasi_ramak_kala_analizi.xlsx")
    a = p.parse_args(argv)
    ornek = a.olaylar is None
    o_yol = ORNEK / "olaylar.csv" if ornek else a.olaylar
    b_yol = ORNEK / "bolumler.csv" if ornek and a.bolumler is None else a.bolumler
    saat, calisan, sc, ac = para(a.saat), para(a.calisan), para(a.siklik_carpan), para(a.agirlik_carpan)
    if (a.saat and (saat is None or saat <= 0)) or (a.calisan and (calisan is None or calisan <= 0)) or sc is None or sc <= 0 or ac is None or ac <= 0 or a.tekrar < 2:
        print("[X] --saat, --calisan ve çarpanlar pozitif sayı; --tekrar en az 2 olmalı")
        return 2
    for y in (o_yol, b_yol):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(o_yol, b_yol, a.cikti, saat=saat, calisan=calisan, siklik_carpan=sc, agirlik_carpan=ac, tekrar=a.tekrar)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {o['olay']} olay · iş kazası {o['kaza']}, ramak kala {o['ramak']} · kayıp gün {o['kayip']:f}" + (f" · sıklık oranı {o['siklik']} · ağırlık oranı {o['agirlik']}" if o["siklik"] is not None else " · çalışma saati verilmedi, oran yok"))
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
