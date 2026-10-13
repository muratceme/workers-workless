import sys
import tempfile
import unittest
from datetime import date, datetime, time
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
K, R = "İş Kazası", "Ramak Kala"


def o(gun, tur=K, bolum="B", neden="N", kayip=None, saat=None, vardiya="", ay=3):
    return {"tarih": date(2026, ay, gun), "saat": saat, "tur": tur, "bolum": bolum, "neden": neden, "kayip": None if kayip is None else D(str(kayip)), "vardiya": vardiya, "aciklama": "", "satir": 2}


class AnalizTesti(unittest.TestCase):
    def test_oranlar(self):
        l = [o(2, kayip=10), o(3, kayip=0), o(4), o(5, R, kayip=99), o(6, R)]
        r = main.analiz(l, D(200000), D(50))
        self.assertEqual(r["ozet"], {"kaza": 3, "ramak": 2, "kayip": D(10), "olay": 5, "kayipli_kaza": 1, "saat": D(200000), "calisan": D(50), "siklik": D("15.00"), "agirlik": D("0.05"), "bin_calisan": D("60.00"),
                                     "ramak_kaza": D("0.67"), "ort_kayip": D("3.33")})
        self.assertEqual([(u["onem"], u["tur"]) for u in r["uyarilar"]], [("Orta", "Tekrarlayan neden"), ("Bilgi", "Kayıp gün yazılmamış")])
        x = main.analiz(l, D(200000), siklik_carpan=D(200000), agirlik_carpan=D(1000000))["ozet"]
        self.assertEqual((x["siklik"], x["agirlik"], x["bin_calisan"], x["calisan"]), (D("3.00"), D("50.00"), None, None))
        r = main.analiz([o(2, R), o(3, R)])
        self.assertEqual((r["ozet"]["siklik"], r["ozet"]["ramak_kaza"], r["ozet"]["ort_kayip"], [u["tur"] for u in r["uyarilar"]]), (None, None, None, ["Çalışma saati verilmedi"]))
        self.assertEqual([u["tur"] for u in main.analiz([o(2, kayip=1)], D(1000))["uyarilar"]], ["Ramak kala kaydı yok"])
        self.assertEqual(main.analiz([o(2, kayip=1)], D(1000))["ozet"]["siklik"], D("1000.00"))
        with self.assertRaises(ValueError):
            main.analiz([])

    def test_bolum_ve_neden(self):
        l = [o(2, bolum="Üretim", neden="Kayma", kayip=6), o(3, bolum="üretim", neden="kayma", kayip=4), o(4, R, "Depo", "Kayma"), o(5, R, "Depo", "Forklift"), o(6, bolum="Depo", neden="Forklift", kayip=2), o(7, R, "Lab", ""),
             o(8, R, "Depo", "Elektrik"), o(9, R, "Depo", "Araç")]
        l[5]["neden"] = "(neden yok)"
        bilgi = [{"bolum": "ÜRETİM", "calisan": D(100), "saat": D(200000)}, {"bolum": "Depo", "calisan": D(20), "saat": D(40000)}, {"bolum": "Ofis", "calisan": D(30), "saat": None}]
        r = main.analiz(l, bolum_bilgi=bilgi, tekrar=2)
        self.assertEqual([(b["ad"], b["kaza"], b["ramak"], b["kayip"], b["calisan"], b["saat"], b["siklik"], b["agirlik"]) for b in r["bolumler"]],
                         [("Üretim", 2, 0, D(10), D(100), D(200000), D("10.00"), D("0.05")), ("Depo", 1, 4, D(2), D(20), D(40000), D("25.00"), D("0.05")), ("Lab", 0, 1, D(0), None, None, None, None),
                          ("Ofis", 0, 0, D(0), D(30), None, None, None)])
        self.assertEqual((r["ozet"]["saat"], r["ozet"]["calisan"], r["ozet"]["siklik"], r["ozet"]["agirlik"], r["ozet"]["bin_calisan"]), (D(240000), D(150), D("12.50"), D("0.05"), D("20.00")))
        self.assertEqual([(n["ad"], n["toplam"], n["kaza"], n["ramak"], n["kayip"], n["pay"], n["kumulatif"]) for n in r["nedenler"]],
                         [("Kayma", 3, 2, 1, D(10), D("37.50"), D("37.50")), ("Forklift", 2, 1, 1, D(2), D("25.00"), D("62.50")), ("Araç", 1, 0, 1, D(0), D("12.50"), D("75.00")), ("Elektrik", 1, 0, 1, D(0), D("12.50"), D("87.50")),
                          ("(neden yok)", 1, 0, 1, D(0), D("12.50"), D("100.00"))])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]), [("Orta", "Bölüm bilgisi yok", "Lab"), ("Orta", "Tekrarlayan neden", "Forklift"), ("Orta", "Tekrarlayan neden", "Kayma")])
        self.assertEqual([u["kim"] for u in main.analiz(l, D(1))["uyarilar"]], ["Kayma"])
        self.assertEqual(main.analiz(l, D(480000), D(10), bilgi)["ozet"]["siklik"], D("6.25"))

    def test_zaman(self):
        l = [o(2, saat=0, vardiya="Gece", kayip=1), o(2, R, saat=3), o(9, saat=4, vardiya="gece"), o(15, R, saat=23, vardiya="Gündüz", ay=4), o(16, saat=12, ay=4), o(1, R, ay=3), o(8, saat=None, ay=3)]
        r = main.analiz(l, D(1))
        k = lambda alan: [(s["ad"], s["kaza"], s["ramak"]) for s in r[alan]]   # noqa: E731
        self.assertEqual(k("aylar"), [("2026-03", 3, 2), ("2026-04", 1, 1)])
        self.assertEqual(k("gunler"), [("Pazartesi", 2, 1), ("Salı", 0, 0), ("Çarşamba", 0, 1), ("Perşembe", 1, 0), ("Cuma", 0, 0), ("Cumartesi", 0, 0), ("Pazar", 1, 1)])
        self.assertEqual(k("saatler"), [("00–04", 1, 1), ("04–08", 1, 0), ("08–12", 0, 0), ("12–16", 1, 0), ("16–20", 0, 0), ("20–24", 0, 1), ("(saat yok)", 1, 1)])
        self.assertEqual((k("vardiyalar"), r["aylar"][0]["kayip"]), ([("Gece", 2, 0), ("Gündüz", 0, 1)], D(1)))
        self.assertEqual(main.analiz([o(2, R)], D(1))["vardiyalar"], [])
        self.assertEqual([main.saat_oku(x) for x in ("09:15", "9.05", "23:59:10", "7", "24:00", "12:60", "", "x", time(14, 30), datetime(2026, 1, 1, 22, 5), None)], [9, 9, 23, 7, None, None, None, None, 14, 22, None])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "olaylar.csv", ORNEK / "bolumler.csv", cikti)
            self.assertEqual(r["ozet"], {"kaza": 7, "ramak": 9, "kayip": D(64), "olay": 16, "kayipli_kaza": 6, "saat": D(480000), "calisan": D(220), "siklik": D("14.58"), "agirlik": D("0.13"), "bin_calisan": D("31.82"),
                                         "ramak_kaza": D("1.29"), "ort_kayip": D("9.14")})
            self.assertEqual([(b["ad"], b["kaza"], b["ramak"], b["kayip"], b["siklik"], b["agirlik"]) for b in r["bolumler"]],
                             [("Üretim", 4, 4, D(45), D("15.15"), D("0.17")), ("Depo", 2, 3, D(19), D("22.73"), D("0.22")), ("Bakım", 1, 1, D(0), D("22.73"), D("0.00")), ("Kalite Laboratuvarı", 0, 1, D(0), None, None),
                              ("Ofis", 0, 0, D(0), D("0.00"), D("0.00"))])
            self.assertEqual([(n["ad"], n["toplam"], n["kaza"], n["kayip"], n["pay"], n["kumulatif"]) for n in r["nedenler"]],
                             [("Kayma ve düşme", 4, 2, D(12), D("25.00"), D("25.00")), ("Forklift ve araç hareketi", 3, 1, D(15), D("18.75"), D("43.75")), ("El aleti kullanımı", 2, 1, D(3), D("12.50"), D("56.25")),
                              ("Elektrik", 2, 1, D(0), D("12.50"), D("68.75")), ("Malzeme düşmesi", 2, 0, D(0), D("12.50"), D("81.25")), ("Elle taşıma", 1, 1, D(4), D("6.25"), D("87.50")),
                              ("Makine koruyucusu eksikliği", 1, 1, D(30), D("6.25"), D("93.75")), ("Kimyasal", 1, 0, D(0), D("6.25"), D("100.00"))])
            k = lambda alan: [(s["ad"], s["kaza"], s["ramak"]) for s in r[alan]]   # noqa: E731
            self.assertEqual(k("aylar"), [(f"2026-{a:02d}", x, y) for a, (x, y) in enumerate(((1, 0), (1, 1), (0, 2), (1, 1), (1, 0), (1, 1), (1, 1), (0, 2), (1, 1)), 1)])
            self.assertEqual(k("gunler"), [("Pazartesi", 1, 1), ("Salı", 0, 3), ("Çarşamba", 2, 2), ("Perşembe", 3, 1), ("Cuma", 1, 0), ("Cumartesi", 0, 2), ("Pazar", 0, 0)])
            self.assertEqual(k("saatler"), [("00–04", 1, 0), ("04–08", 0, 0), ("08–12", 2, 5), ("12–16", 1, 3), ("16–20", 1, 1), ("20–24", 2, 0)])
            self.assertEqual(k("vardiyalar"), [("Gece", 3, 0), ("Gündüz", 4, 9)])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Orta", "Bölüm bilgisi yok", "Kalite Laboratuvarı"), ("Orta", "Tekrarlayan neden", "Forklift ve araç hareketi"), ("Orta", "Tekrarlayan neden", "Kayma ve düşme")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Bölümler", "Neden Pareto", "Aylar", "Haftanın Günü", "Saat Dilimi", "Vardiya", "Olaylar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            oy, by = Path(tmp) / "o.csv", Path(tmp) / "b.csv"
            oy.write_text("Kaza Tarihi;Olay Saati;Olay Türü;Birim;Kök Neden;Rapor Günü\n05.03.2026;14.30;KAZA;;;2,5\n06.03.2026;;ramak kala;Depo;Kayma;\nx;;Kaza;;;\n07.03.2026;;Meslek Hastalığı;;;\n08.03.2026;;Kaza;;;-1\n;;;;;\n",
                          encoding="utf-8")
            by.write_text("Departman;Personel;Toplam Çalışma Saati\nDepo;12;26.400\n;5;5\n", encoding="utf-8")
            l, uy = main.oku_olaylar(oy)
            self.assertEqual(([(x["tarih"], x["saat"], x["tur"], x["bolum"], x["neden"], x["kayip"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([(date(2026, 3, 5), 14, K, "(bölüm yok)", "(neden yok)", D("2.5")), (date(2026, 3, 6), None, R, "Depo", "Kayma", None)], [("Olay okunamadı", f"Satır {n}") for n in (4, 5, 6)]))
            self.assertEqual(main.oku_bolumler(by), [{"bolum": "Depo", "calisan": D(12), "saat": D(26400)}])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--olaylar", str(oy), "--cikti", c]
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(load_workbook(c).sheetnames, ["Özet", "Bölümler", "Neden Pareto", "Aylar", "Haftanın Günü", "Saat Dilimi", "Olaylar", "Uyarılar"])
            self.assertEqual(main.main(arg + ["--bolumler", str(by), "--saat", "480.000", "--calisan", "220", "--siklik-carpan", "200000", "--agirlik-carpan", "1000000", "--tekrar", "2"]), 0)
            self.assertEqual([main.main(arg + x) for x in (["--saat", "0"], ["--saat", "x"], ["--calisan", "0"], ["--siklik-carpan", "0"], ["--agirlik-carpan", "x"], ["--tekrar", "1"])], [2] * 6)
            self.assertEqual(main.main(arg + ["--bolumler", str(Path(tmp) / "yok.csv")]), 1)
            by.write_text("Bölüm;Çalışan Sayısı\nA;x\n", encoding="utf-8")
            self.assertEqual(main.main(arg + ["--bolumler", str(by)]), 1)
            for icerik in ("Tarih;Bölüm\n05.03.2026;A\n", "Tarih;Tür\n;\n"):
                oy.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
