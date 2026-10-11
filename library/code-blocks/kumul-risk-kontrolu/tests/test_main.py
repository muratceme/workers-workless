import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
RAPOR = date(2026, 10, 1)
M = 1_000_000


def pol(no, net, bolge="İstanbul", lokasyon=""):
    return {"no": no, "sigortali": "", "bolge": bolge, "il": "", "ilce": "", "adres": lokasyon, "lokasyon": ("A:" + main.katla(lokasyon)) if lokasyon else "", "bedel": D(net), "pay": D(100),
            "net": D(net)}


KAP = {"istanbul": {"kapasite": D(400), "pml": D(15)}, "izmir": {"kapasite": D(100), "pml": D(100)}}


class KumulTesti(unittest.TestCase):
    def test_bolge_kumulu(self):
        r = main.kumul([pol("P1", 1000), pol("P2", 1300), pol("P3", 100, "İzmir"), pol("P4", 50, "Antalya")], KAP)
        d = {b["ad"]: (b["police"], b["bedel"], b["pml"], b["kullanim"], b["bos"], b["durum"]) for b in r["bolge"]}
        self.assertEqual(d["İstanbul"], (2, D(2300), D("345.00"), D("86.25"), D("55.00"), "Eşiğe yakın"))
        self.assertEqual(d["İzmir"], (1, D(100), D("100.00"), D("100.00"), D("0.00"), "Eşiğe yakın"))          # tam kapasite aşım değildir
        self.assertEqual(d["Antalya"], (1, D(50), None, None, None, "Kapasite tanımsız"))
        self.assertEqual(([u["tur"] for u in r["uyarilar"]], r["toplam"]), (["Bölge kapasitesi tanımsız"], D(2450)))
        self.assertEqual(main.kumul([pol("P1", 101, "İzmir")], KAP)["bolge"][0]["durum"], "Aşım")
        self.assertEqual([main.kumul([pol("P1", n, "İzmir")], KAP, uyari_esik=D(90))["bolge"][0]["durum"] for n in (89, 90)], ["Uygun", "Eşiğe yakın"])
        self.assertEqual([b["ad"] for b in r["bolge"]], ["İzmir", "İstanbul", "Antalya"])

    def test_lokasyon_kumulu(self):
        l = [pol("P1", 900, lokasyon="OSB 4. Cadde No: 12"), pol("P2", 400, lokasyon="osb 4 cadde no:12"), pol("P3", 1200, lokasyon="GOSB 7"), pol("P4", 800, lokasyon="X Sokak 1"),
             pol("P5", 790, lokasyon="Y Sokak 2"), pol("P6", 10, lokasyon="Z"), pol("P7", 20, lokasyon="Z"), pol("P8", 5000)]
        r = main.kumul(l, KAP, D(1000))
        self.assertEqual([(len(x["policeler"]), x["bedel"], x["kullanim"], x["asim"], x["durum"]) for x in r["lokasyon"]],
                         [(2, D(1300), D("130.00"), D(300), "Aşım"), (1, D(1200), D("120.00"), D(200), "Aşım"), (1, D(800), D("80.00"), D(0), "Eşiğe yakın"), (2, D(30), D("3.00"), D(0), "Uygun")])
        r = main.kumul(l, KAP)                                                                                   # kapasite verilmezse yalnız çok poliçeli lokasyonlar
        self.assertEqual([(len(x["policeler"]), x["durum"], x["kullanim"]) for x in r["lokasyon"]], [(2, "Birden çok poliçe", None)] * 2)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "k.xlsx"
            r = main.calistir(ORNEK / "policeler.csv", ORNEK / "kapasiteler.csv", cikti, RAPOR, D(1000 * M))
            self.assertEqual((len(r["policeler"]), r["disarida"], r["toplam"]), (10, {"bitmis": 1, "baslamamis": 1}, D(5450 * M)))
            d = {b["ad"]: (b["bedel"], b["pml"], b["kullanim"], b["durum"]) for b in r["bolge"]}
            self.assertEqual(d["İstanbul"], (D(2300 * M), D(345 * M), D("86.25"), "Eşiğe yakın"))                 # K-03 yüzde 50 paylı: 300 milyon
            self.assertEqual(d["İzmir"], (D(1250 * M), D(150 * M), D("107.14"), "Aşım"))
            self.assertEqual(d["Bursa-Kocaeli"], (D(1500 * M), D(225 * M), D("75.00"), "Uygun"))
            self.assertEqual((d["Diğer"][2:], d["Antalya"][3]), ((D("12.50"), "Uygun"), "Kapasite tanımsız"))
            self.assertEqual([([p["no"] for p in x["policeler"]], x["bedel"], x["durum"]) for x in r["lokasyon"]],
                             [(["K-01", "K-02"], D(1300 * M), "Aşım"), (["K-08"], D(1200 * M), "Aşım"), (["K-06"], D(800 * M), "Eşiğe yakın")])
            self.assertEqual([u["tur"] for u in r["uyarilar"]], ["Bölge kapasitesi tanımsız"])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Bölge Kümülü", "Lokasyon Kümülü", "Poliçeler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            p, k = Path(tmp) / "p.csv", Path(tmp) / "k.csv"
            p.write_text("Poliçe No;Bölge;Adres;Adres Kodu;Bedel;Şirket Payı %;Başlangıç;Bitiş\nP1;A;X Sokak;1234567;1.000,00;40;01.10.2026;01.10.2027\nP1;A;;;5;;;\nP2;A;;;çok;;;\n"
                         "P3;;;;100;150;bozuk;\nP4;A;Y;;10;;01.10.2025;01.10.2026\nP5;A;Y;;10;;02.10.2026;02.10.2027\n", encoding="utf-8")
            l, dis, uy = main.oku(p, RAPOR)
            self.assertEqual([(x["no"], x["net"], x["lokasyon"], x["bolge"]) for x in l], [("P1", D("400.00"), "K:1234567", "A"), ("P3", D("100.00"), "", "Belirtilmemiş")])
            self.assertEqual((dis, sorted(u["tur"] for u in uy)),
                             ({"bitmis": 1, "baslamamis": 1}, ["Adres bilgisi yok", "Bedel okunamadı", "Kümül bölgesi boş", "Mükerrer poliçe", "Tarih okunamadı", "Şirket payı okunamadı"]))
            k.write_text("Bölge;Kapasite;PML %\nA;500;20\nB;yok;\nC;100;120\n", encoding="utf-8")
            kap, uy = main.kapasite_oku(k)
            self.assertEqual((kap, len(uy)), ({"a": {"kapasite": D(500), "pml": D(20)}, "c": {"kapasite": D(100), "pml": D(100)}}, 2))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--policeler", str(p), "--kapasiteler", str(k), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "01.10.2026", "--lokasyon-kapasite", "300"]), 0)
            self.assertEqual(main.main(arg + ["--tarih", "dün"]), 2)
            self.assertEqual(main.main(arg + ["--lokasyon-kapasite", "0"]), 2)
            p.write_text("Poliçe No;Bedel;Başlangıç;Bitiş\nP1;100;01.10.2025;01.10.2026\n", encoding="utf-8")
            self.assertEqual(main.main(arg + ["--tarih", "01.10.2026"]), 1)                                        # yürürlükte poliçe yok
            self.assertEqual(main.main(["--policeler", str(Path(tmp) / "yok.csv")]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
