"""
OEE Hesaplama — Workers / Workless kod bloğu
Üretim › Üretim Mühendisi

Vardiya/makine bazında planlı süre, duruşlar, ideal çevrim süresi ve üretim adetlerinden
Kullanılabilirlik, Performans, Kalite ve OEE'yi hesaplar; kayıpları dakika olarak ayrıştırır ve
makine bazında (süre ve adet toplamlarıyla, oran ortalamasıyla değil) özetler. İnternete bağlanmaz.

  Planlı üretim süresi = vardiya süresi − planlı duruş (mola, planlı bakım)
  Çalışma süresi       = planlı üretim süresi − plansız duruş (arıza, ayar, malzeme bekleme)
  Kullanılabilirlik    = çalışma süresi / planlı üretim süresi
  Performans           = ideal çevrim süresi × toplam adet / çalışma süresi
  Kalite               = sağlam adet / toplam adet
  OEE                  = Kullanılabilirlik × Performans × Kalite

Kullanım:
    python main.py                                  # örnek vardiya kayıtlarıyla dener
    python main.py --girdi vardiya_kayitlari.xlsx
"""
from __future__ import annotations

import argparse
import csv
import sys
from collections import OrderedDict
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent


def sayi(x) -> float | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return float(x)
    try:
        return float(str(x).strip().replace(".", "").replace(",", ".") if "," in str(x) else str(x).strip())
    except ValueError:
        return None


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


SUTUNLAR = {
    "tarih": ("tarih",), "vardiya": ("vardiya",), "makine": ("makine", "hat", "iş merkezi", "tezgah"),
    "urun": ("ürün", "ürün kodu", "malzeme"),
    "vardiya_suresi": ("vardiya süresi (dk)", "vardiya süresi", "toplam süre (dk)", "süre (dk)"),
    "planli_durus": ("planlı duruş (dk)", "planlı duruş", "mola (dk)"),
    "plansiz_durus": ("plansız duruş (dk)", "plansız duruş", "arıza (dk)", "duruş (dk)"),
    "cevrim": ("ideal çevrim süresi (sn)", "ideal çevrim (sn)", "çevrim süresi (sn)", "ideal çevrim süresi"),
    "toplam": ("toplam üretim", "toplam adet", "üretilen", "üretim adedi"),
    "hatali": ("hatalı", "hatalı adet", "fire", "hurda", "ıskarta"),
}
ZORUNLU = ("makine", "vardiya_suresi", "plansiz_durus", "cevrim", "toplam", "hatali")


def liste_oku(yol: Path) -> list[dict]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ayirici = csv.Sniffer().sniff(metin.splitlines()[0], delimiters=",;\t").delimiter
        satirlar = list(csv.reader(metin.splitlines(), delimiter=ayirici))
    satirlar = [r for r in satirlar if any(c not in (None, "") for c in r)]
    b = [kucuk(x) for x in satirlar[0]]
    k = {alan: next((i for i, x in enumerate(b) if x in adlar), None) for alan, adlar in SUTUNLAR.items()}
    eksik = [a for a in ZORUNLU if k[a] is None]
    if eksik:
        raise SystemExit(f"Gerekli sütunlar eksik: {eksik}. Başlıklar: {satirlar[0]}")
    return [{alan: (r[i] if i is not None and i < len(r) else None) for alan, i in k.items()} for r in satirlar[1:]]


@dataclass
class Olcum:
    planli_sure: float      # dk
    calisma: float          # dk
    ideal_sure: float       # dk (ideal çevrim × toplam)
    toplam: float
    saglam: float

    @property
    def kullanilabilirlik(self):
        return self.calisma / self.planli_sure if self.planli_sure else 0.0

    @property
    def performans(self):
        return self.ideal_sure / self.calisma if self.calisma else 0.0

    @property
    def kalite(self):
        return self.saglam / self.toplam if self.toplam else 0.0

    @property
    def oee(self):
        return self.kullanilabilirlik * self.performans * self.kalite

    def __add__(self, o: "Olcum") -> "Olcum":
        return Olcum(self.planli_sure + o.planli_sure, self.calisma + o.calisma, self.ideal_sure + o.ideal_sure,
                     self.toplam + o.toplam, self.saglam + o.saglam)

    def kayiplar(self, cevrim_dk_ort: float) -> dict:
        kalite_kaybi = (self.toplam - self.saglam) * cevrim_dk_ort
        return {
            "Kullanılabilirlik kaybı (dk)": self.planli_sure - self.calisma,
            "Performans kaybı (dk)": self.calisma - self.ideal_sure,
            "Kalite kaybı (dk)": kalite_kaybi,
            "Değer katan süre (dk)": self.ideal_sure - kalite_kaybi,
        }


def satir_olcum(k: dict) -> tuple[Olcum | None, list[str]]:
    uyarilar = []
    vs, pd, ud = sayi(k["vardiya_suresi"]), sayi(k.get("planli_durus")) or 0.0, sayi(k["plansiz_durus"]) or 0.0
    cevrim, toplam, hatali = sayi(k["cevrim"]), sayi(k["toplam"]), sayi(k["hatali"]) or 0.0
    if None in (vs, cevrim, toplam):
        return None, ["Vardiya süresi, ideal çevrim süresi veya toplam üretim eksik"]
    planli = vs - pd
    calisma = planli - ud
    if calisma <= 0:
        return None, ["Çalışma süresi sıfır veya negatif (duruşlar vardiya süresini aşıyor)"]
    if hatali > toplam:
        uyarilar.append("Hatalı adet toplam üretimden büyük")
    ideal = cevrim * toplam / 60
    o = Olcum(planli, calisma, ideal, toplam, max(toplam - hatali, 0))
    if o.performans > 1.0:
        uyarilar.append(f"Performans %{o.performans * 100:.1f}: ideal çevrim süresi gerçekçi değil veya adet/süre hatalı")
    return o, uyarilar


BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
YUZDE = "0.0%"


def renk(oee: float) -> PatternFill:
    return PatternFill("solid", fgColor="B7E4C7" if oee >= 0.85 else "FFE08A" if oee >= 0.60 else "FDE2E1")


def calistir(girdi: Path, cikti: Path) -> dict:
    kayitlar = liste_oku(girdi)
    wb = Workbook()
    ws = wb.active
    ws.title = "Vardiya OEE"
    basliklar = ["Tarih", "Vardiya", "Makine", "Ürün", "Planlı Süre (dk)", "Çalışma (dk)", "Toplam", "Sağlam",
                 "Kullanılabilirlik", "Performans", "Kalite", "OEE", "Uyarılar"]
    ws.append(basliklar)
    for h in ws[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    makineler: "OrderedDict[str, list]" = OrderedDict()
    for k in kayitlar:
        o, uyarilar = satir_olcum(k)
        if o is None:
            ws.append([k.get("tarih"), k.get("vardiya"), k.get("makine"), k.get("urun")] + [None] * 8 + ["; ".join(uyarilar)])
            continue
        ws.append([k.get("tarih"), k.get("vardiya"), k.get("makine"), k.get("urun"), o.planli_sure, o.calisma, o.toplam,
                   o.saglam, o.kullanilabilirlik, o.performans, o.kalite, o.oee, "; ".join(uyarilar)])
        for c in (9, 10, 11, 12):
            ws.cell(ws.max_row, c).number_format = YUZDE
        ws.cell(ws.max_row, 12).fill = renk(o.oee)
        m = makineler.setdefault(str(k["makine"]), [Olcum(0, 0, 0, 0, 0), 0.0, 0.0])
        m[0] = m[0] + o
        m[1] += (sayi(k["cevrim"]) or 0) * o.toplam / 60        # ağırlıklı çevrim (dk) için pay
        m[2] += o.toplam
    for j, g in enumerate((12, 9, 12, 12, 13, 12, 9, 9, 13, 12, 9, 9, 60), 1):
        ws.column_dimensions[get_column_letter(j)].width = g
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions

    oz = wb.create_sheet("Makine Özeti")
    oz.append(["Makine", "Kullanılabilirlik", "Performans", "Kalite", "OEE", "Kullanılabilirlik kaybı (dk)",
               "Performans kaybı (dk)", "Kalite kaybı (dk)", "Değer katan süre (dk)"])
    for h in oz[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    ozet = {}
    genel = Olcum(0, 0, 0, 0, 0)
    for ad, (o, pay, adet) in makineler.items():
        kay = o.kayiplar(pay / adet if adet else 0)
        oz.append([ad, o.kullanilabilirlik, o.performans, o.kalite, o.oee, *[round(v, 1) for v in kay.values()]])
        for c in (2, 3, 4, 5):
            oz.cell(oz.max_row, c).number_format = YUZDE
        oz.cell(oz.max_row, 5).fill = renk(o.oee)
        ozet[ad] = o
        genel = genel + o
    if makineler:
        oz.append(["TOPLAM", genel.kullanilabilirlik, genel.performans, genel.kalite, genel.oee])
        for c in (2, 3, 4, 5):
            oz.cell(oz.max_row, c).number_format = YUZDE
        for h in oz[oz.max_row]:
            h.font = Font(bold=True)
        g = BarChart()
        g.type, g.title, g.height, g.width = "col", "Makine bazında OEE bileşenleri", 9, 20
        g.add_data(Reference(oz, min_col=2, max_col=5, min_row=1, max_row=1 + len(makineler)), titles_from_data=True)
        g.set_categories(Reference(oz, min_col=1, min_row=2, max_row=1 + len(makineler)))
        g.y_axis.number_format = "0%"
        oz.add_chart(g, "B" + str(oz.max_row + 3))
    for j, gen in enumerate((14, 16, 13, 10, 10, 24, 22, 18, 22), 1):
        oz.column_dimensions[get_column_letter(j)].width = gen

    bilgi = wb.create_sheet("Bilgi")
    for s in [["Kullanılabilirlik", "Çalışma süresi ÷ planlı üretim süresi (vardiya − planlı duruş)"],
              ["Performans", "İdeal çevrim süresi × toplam adet ÷ çalışma süresi"],
              ["Kalite", "Sağlam adet ÷ toplam adet"], ["OEE", "Kullanılabilirlik × Performans × Kalite"],
              ["Toplama", "Makine ve genel değerler süre/adet toplamlarından hesaplanır (oranların ortalaması değil)"],
              ["Renk", "Yeşil ≥ %85 (sıklıkla 'dünya standardı' olarak anılan düzey) · Sarı ≥ %60 · Kırmızı < %60"]]:
        bilgi.append(s)
    bilgi.column_dimensions["A"].width = 18
    bilgi.column_dimensions["B"].width = 100
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    return {"makineler": ozet, "genel": genel}


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description="Vardiya/makine bazında OEE hesabı.")
    ap.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "vardiya_kayitlari.csv")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "oee.xlsx")
    a = ap.parse_args(argv)
    s = calistir(a.girdi, a.cikti)
    for ad, o in s["makineler"].items():
        print(f"[OK] {ad}: OEE %{o.oee * 100:.1f} (K %{o.kullanilabilirlik * 100:.1f} · P %{o.performans * 100:.1f} · Q %{o.kalite * 100:.1f})".replace(".", ","))
    print(f"[OK] Genel OEE %{s['genel'].oee * 100:.1f}".replace(".", ","))
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
