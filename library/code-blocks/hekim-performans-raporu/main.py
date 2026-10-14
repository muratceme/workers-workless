"""
Hekim Performans Raporu — Workers / Workless kod bloğu
Hastane Yönetimi › Hastane Müdürü
Hekim bazında muayene, ameliyat, yatış ve gelir göstergelerini raporlar; her hekimi kendi bölümünün ortalamasıyla karşılaştırır.
  Çalışma günü        = Çalışma Günü sütunu; yoksa hekimin kaydı olan farklı gün sayısı
  Günlük muayene      = muayene ÷ çalışma günü (ameliyat ve gelir için de aynı)
  Muayeneden yatış %  = yatış ÷ muayene × 100
  Bölüme göre endeks  = hekimin günlük değeri ÷ bölümün günlük değeri × 100 (bölüm değeri toplamların oranıdır)
  Bilgi notu          = endeks 100 ± --sapma (50) dışında — inceleme önerisidir, değerlendirme değildir
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"hekim": ("hekim", "doktor", "hekim adi", "uzman"), "bolum": ("bolum", "brans", "klinik"), "tarih": ("tarih", "gun"), "muayene": ("muayene", "poliklinik", "muayene sayisi"),
         "ameliyat": ("ameliyat", "ameliyat sayisi", "operasyon"), "yatis": ("yatis", "yatis sayisi", "yatirilan"), "gelir": ("gelir", "ciro", "faturalanan"),
         "calisma": ("calisma gunu", "calisilan gun", "fiili gun")}
SAYI = ("muayene", "ameliyat", "yatis", "gelir", "calisma")
SIFIR = Decimal(0)
OLCU = (("muayene", "Günlük muayene"), ("ameliyat", "Günlük ameliyat"), ("gelir", "Günlük gelir"))


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("hekim",)):
        hekim = metin(r["hekim"])
        if not hekim:
            continue
        d = {a: para(r[a]) if metin(r[a]) else None for a in SAYI}
        t = tarih(r["tarih"]) if metin(r["tarih"]) else None
        if any(metin(r[a]) and (d[a] is None or d[a] < 0) for a in SAYI) or (metin(r["tarih"]) and t is None):
            uy.append(uyari("Yüksek", "Satır okunamadı", hekim, f"Satır {r['_satir']}: tarih ya da sayılardan biri okunamadı. Satır alınmadı."))
            continue
        l.append({"hekim": hekim, "bolum": metin(r["bolum"]) or "Belirtilmemiş", "tarih": t, **d})
    return l, uy


def _ozet(l: list[dict]) -> dict:
    top = lambda a: sum((x[a] or SIFIR for x in l), SIFIR)   # noqa: E731
    yazili = [x["calisma"] for x in l if x["calisma"] is not None]
    gun = sum(yazili, SIFIR) if yazili else Decimal(len({x["tarih"] for x in l if x["tarih"]}))
    o = {"gun": gun, "muayene": top("muayene"), "ameliyat": top("ameliyat"), "yatis": top("yatis"), "gelir": top("gelir")}
    for a, _ in OLCU:
        o[a + "_gun"] = yuvarla(o[a] / gun) if gun else None
    o["yatis_oran"] = yuvarla(o["yatis"] * 100 / o["muayene"]) if o["muayene"] else None
    return o


def rapor(kayit: list[dict], sapma: Decimal = Decimal(50)) -> dict:
    if not kayit:
        raise ValueError("Hekim kaydı yok.")
    uy, bolumler, sonuc = [], {}, []
    for x in kayit:
        bolumler.setdefault(katla(x["bolum"]), []).append(x)
    ozet = []
    for k in sorted(bolumler, key=lambda k: katla(bolumler[k][0]["bolum"])):
        l = bolumler[k]
        b = {"bolum": l[0]["bolum"], **_ozet(l)}
        hekimler = {}
        for x in l:
            hekimler.setdefault(katla(x["hekim"]), []).append(x)
        b["hekim_sayisi"] = len(hekimler)
        if not b["gun"]:
            uy.append(uyari("Orta", "Çalışma günü yok", b["bolum"], "Ne Çalışma Günü ne Tarih yazılmış; günlük değerler ve endeksler hesaplanmadı."))
        satirlar = []
        for hl in hekimler.values():
            h = {"hekim": hl[0]["hekim"], "bolum": b["bolum"], **_ozet(hl)}
            h["gelir_pay"] = yuvarla(h["gelir"] * 100 / b["gelir"]) if b["gelir"] else None
            for a, ad in OLCU:
                h[a + "_endeks"] = yuvarla(h[a + "_gun"] * 100 / b[a + "_gun"]) if h[a + "_gun"] is not None and b[a + "_gun"] else None
                e = h[a + "_endeks"]
                if e is not None and len(hekimler) > 1 and (e < 100 - sapma or e > 100 + sapma):
                    uy.append(uyari("Bilgi", f"{ad} bölüm ortalamasından {'düşük' if e < 100 else 'yüksek'}", f"{b['bolum']} · {h['hekim']}",
                                    f"{ad} {h[a + '_gun']:,.2f}, bölüm {b[a + '_gun']:,.2f} (endeks {e}). Vaka karmaşıklığı, nöbet, idari görev ve izinleri göz önüne alın."))
            satirlar.append(h)
        sonuc += sorted(satirlar, key=lambda h: (-h["gelir"], katla(h["hekim"])))
        ozet.append(b)
    return {"hekimler": sonuc, "bolumler": ozet, "toplam": _ozet(kayit), "uyarilar": uy}


BASLIK = ["Çalışma Günü", "Muayene", "Ameliyat", "Yatış", "Muayeneden Yatış %", "Gelir", "Günlük Muayene", "Günlük Ameliyat", "Günlük Gelir"]


def _deger(o: dict) -> list:
    return [o["gun"], o["muayene"], o["ameliyat"], o["yatis"], o["yatis_oran"], o["gelir"], o["muayene_gun"], o["ameliyat_gun"], o["gelir_gun"]]


def rapor_yaz(cikti: Path, r: dict, sapma: Decimal) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Hekimler"
    _baslik(ws, ["Bölüm", "Hekim", "Gelir Payı %"] + BASLIK + ["Muayene Endeksi", "Ameliyat Endeksi", "Gelir Endeksi"], (18, 22, 9, 8, 9, 9, 8, 10, 14, 10, 10, 13, 10, 10, 10), "C2")
    renk = lambda e: None if e is None else SARI if e < 100 - sapma or e > 100 + sapma else None   # noqa: E731
    for h in r["hekimler"]:
        _satir(ws, [h["bolum"], h["hekim"], h["gelir_pay"]] + _deger(h) + [h["muayene_endeks"], h["ameliyat_endeks"], h["gelir_endeks"]], para_s=(3, 8, 9, 10, 11, 12, 13, 14, 15),
               dolgu={13: renk(h["muayene_endeks"]), 14: renk(h["ameliyat_endeks"]), 15: renk(h["gelir_endeks"])})
    ws.append([])
    ws.append(["Endeks = hekimin günlük değeri ÷ bölümün günlük değeri × 100. Rapor tek başına performans değerlendirmesi değildir; vaka karmaşıklığı, nöbet ve idari görevler gözetilmez."])
    ws = wb.create_sheet("Bölümler")
    _baslik(ws, ["Bölüm", "Hekim Sayısı"] + BASLIK, (20, 8, 8, 9, 9, 8, 10, 14, 10, 10, 13))
    for b in r["bolumler"] + [{"bolum": "HASTANE", "hekim_sayisi": len({katla(h["hekim"]) for h in r["hekimler"]}), **r["toplam"]}]:
        k = _satir(ws, [b["bolum"], b["hekim_sayisi"]] + _deger(b), para_s=(7, 8, 9, 10, 11))
        if b["bolum"] == "HASTANE":
            for c in ws[k]:
                c.font = Font(bold=True)
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, sapma: Decimal = Decimal(50)) -> dict:
    kayit, uy = oku(yol)
    r = rapor(kayit, sapma)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, sapma)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Hekim bazında muayene, ameliyat, yatış ve gelir göstergelerini raporlar.")
    p.add_argument("--veriler", type=Path, help="Hekim, Bölüm, Tarih (ya da Çalışma Günü), Muayene, Ameliyat, Yatış, Gelir")
    p.add_argument("--sapma", default="50", help="Bölüm ortalamasından bu yüzdeden fazla sapan değerler için bilgi notu (varsayılan 50)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "hekim_performans_raporu.xlsx")
    a = p.parse_args(argv)
    yol = a.veriler or ORNEK / "hekim_aylik_verileri.csv"
    sapma = para(a.sapma)
    if sapma is None or not 0 <= sapma <= 100:
        print("[X] --sapma 0-100 arası olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, sapma)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {len(r['hekimler'])} hekim · {len(r['bolumler'])} bölüm · muayene {t['muayene']:,.0f} · ameliyat {t['ameliyat']:,.0f} · yatış {t['yatis']:,.0f} · gelir {t['gelir']:,.2f}")
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
