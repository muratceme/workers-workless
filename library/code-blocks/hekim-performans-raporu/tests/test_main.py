import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def x(hekim, bolum, muayene=0, ameliyat=0, yatis=0, gelir=0, gun=None, t=None):
    return {"hekim": hekim, "bolum": bolum, "tarih": date(2026, 10, t) if t else None, "muayene": D(muayene), "ameliyat": D(ameliyat), "yatis": D(yatis), "gelir": D(gelir),
            "calisma": None if gun is None else D(gun)}


def veri():
    return [x("H1", "B1", 400, 10, 20, 200000, 20), x("H2", "B1", 100, 10, 5, 100000, 10), x("H3", "B2", 30, t=1), x("h3", "b2", 20, t=2), x("H4", "B3", 5)]


class RaporTesti(unittest.TestCase):
    def test_endeks(self):
        r = main.rapor(veri())
        self.assertEqual([(h["hekim"], h["gun"], h["muayene_gun"], h["muayene_endeks"], h["ameliyat_endeks"], h["gelir_endeks"], h["gelir_pay"]) for h in r["hekimler"]],
                         [("H1", D(20), D("20.00"), D("119.98"), D("74.63"), D("100.00"), D("66.67")), ("H2", D(10), D("10.00"), D("59.99"), D("149.25"), D("100.00"), D("33.33")),
                          ("H3", D(2), D("25.00"), D("100.00"), None, None, None), ("H4", D(0), None, None, None, None, None)])
        b = r["bolumler"][0]
        self.assertEqual((b["bolum"], b["hekim_sayisi"], b["gun"], b["muayene_gun"], b["ameliyat_gun"], b["gelir_gun"], b["yatis_oran"]), ("B1", 2, D(30), D("16.67"), D("0.67"), D("10000.00"), D("5.00")))
        self.assertEqual([(u["tur"], u["kim"]) for u in r["uyarilar"]], [("Çalışma günü yok", "B3")])

    def test_sapma(self):
        r = main.rapor(veri(), D(40))
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"] if u["onem"] == "Bilgi"),
                         [("Günlük ameliyat bölüm ortalamasından yüksek", "B1 · H2"), ("Günlük muayene bölüm ortalamasından düşük", "B1 · H2")])
        with self.assertRaises(ValueError):
            main.rapor([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "hekim_aylik_verileri.csv", cikti)
            self.assertEqual([(h["hekim"], h["muayene_endeks"], h["ameliyat_endeks"], h["gelir_endeks"]) for h in r["hekimler"]][:6],
                             [("Dr. M. T.", D("120.23"), None, D("109.13")), ("Dr. N. Ö.", D("117.09"), None, D("109.89")), ("Dr. H. A.", D("59.71"), None, D("79.49")),
                              ("Dr. A. K.", D("95.91"), D("103.66"), D("103.10")), ("Dr. S. Y.", D("124.99"), D("64.02"), D("79.52")), ("Dr. E. B.", D("63.10"), D("157.32"), D("130.68"))])
            self.assertEqual([(u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Günlük ameliyat bölüm ortalamasından düşük", "Ortopedi · Dr. B. Ç."), ("Günlük ameliyat bölüm ortalamasından yüksek", "Genel Cerrahi · Dr. E. B.")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Hekimler", "Bölümler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "v.csv"
            y.write_text("Doktor;Branş;Poliklinik;Ciro;Fiili Gün\nA;X;1.200;150.000,50;20\nB;X;x;;\nC;X;5;;-1\n;X;5;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(z["hekim"], z["muayene"], z["gelir"], z["calisma"]) for z in l], len(uy)), ([("A", D(1200), D("150000.50"), D(20))], 2))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--veriler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--sapma", "30"]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--sapma", "x"], ["--sapma", "101"])], [2] * 2)
            self.assertEqual(main.main(["--veriler", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Branş;Ciro\nX;5\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
