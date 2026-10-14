import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BAS, BIT = date(2026, 9, 1), date(2026, 9, 10)
SERVIS = {"s1": {"servis": "S1", "yatak": 2}, "s2": {"servis": "S2", "yatak": 1}, "s4": {"servis": "S4", "yatak": 3}}


def y(p, servis, yatis, taburcu=None, plan=None):
    g = lambda x: None if x is None else (date(2026, 8, x[1]) if isinstance(x, tuple) else date(2026, 9, x))   # noqa: E731
    return {"protokol": p, "hasta": "", "servis": servis, "yatis": g(yatis), "taburcu": g(taburcu), "plan": g(plan), "hekim": ""}


def veri():
    return [y("P1", "S1", (8, 28), 3), y("P2", "s1", 2, None, 11), y("P3", "S1", 5, 5), y("P4", "S2", 1, 10), y("P5", "S2", 9, None, 10), y("P6", "S3", 5)]


def k(o):
    return (o["servis"], o["yatak"], o["hasta_gunu"], o["doluluk"], o["kabul"], o["taburcu"], o["alos"], o["taburcu_ort"], o["devir"], o["dolu"], o["bos"])


class RaporTesti(unittest.TestCase):
    def test_doluluk(self):
        r = main.rapor(SERVIS, veri(), BAS, BIT, BIT)
        self.assertEqual([k(o) for o in r["servisler"]],
                         [("S1", 2, 12, D("60.00"), 2, 2, D("6.00"), D("3.50"), D("1.00"), 1, 1), ("S2", 1, 11, D("110.00"), 2, 1, D("11.00"), D("9.00"), D("1.00"), 1, 0),
                          ("S3", None, 6, None, 1, 0, None, None, None, 1, None), ("S4", 3, 0, D(0), 0, 0, None, None, D(0), 0, 3)])
        self.assertEqual(k(r["toplam"]), ("TOPLAM", 6, 23, D("38.33"), 4, 3, D("7.67"), D("5.33"), D("0.50"), 2, 4))
        self.assertEqual(([(a["protokol"], a["gun"]) for a in r["anlik"]], [(a["protokol"], a["gecmis"]) for a in r["beklenen"]]), ([("P2", 8), ("P5", 1), ("P6", 5)], [("P5", False), ("P2", False)]))
        self.assertEqual(sorted((u["tur"], u["kim"]) for u in r["uyarilar"]), [("Doluluk yüksek", "S2"), ("Servis tanımlı değil", "S3"), ("Yatış kaydı yok", "S4")])

    def test_kurallar(self):
        self.assertEqual([main.hasta_gunu(date(2026, 9, a), date(2026, 9, b) if b else None, BAS, BIT) for a, b in ((5, 5), (5, 6), (10, None), (10, 10), (11, 11), (9, 12))], [1, 1, 1, 1, 0, 2])
        r = main.rapor(SERVIS, veri() + [y("P7", "S2", 2, None, 8)], BAS, BIT, BIT, ileri=0, uzun=7, dusuk=D(70))
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Bilgi", "Doluluk düşük", "S1"), ("Bilgi", "Uzun yatış", "S2 · P7"), ("Bilgi", "Uzun yatış", "s1 · P2"), ("Bilgi", "Yatış kaydı yok", "S4"), ("Orta", "Doluluk yüksek", "S2"), ("Orta", "Planlanan taburcu geçti", "S2 · P7"),
                          ("Orta", "Servis tanımlı değil", "S3"), ("Yüksek", "Yatan hasta yatak sayısından fazla", "S2")])
        self.assertEqual([(a["protokol"], a["gecmis"]) for a in r["beklenen"]], [("P7", True), ("P5", False)])
        self.assertIsNone(main.rapor({}, veri(), BAS, BIT, BIT)["toplam"])
        for a, b in (([], BIT), (veri(), date(2026, 8, 1))):
            with self.assertRaises(ValueError):
                main.rapor(SERVIS, a, BAS, b, BIT)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            g = date(2026, 9, 30)
            r = main.calistir(ORNEK / "servisler.csv", ORNEK / "yatislar.csv", cikti, date(2026, 9, 1), g, g)
            self.assertEqual([k(o) for o in r["servisler"]] + [k(r["toplam"])],
                             [("Dahiliye", 12, 312, D("86.67"), 52, 52, D("6.00"), D("5.46"), D("4.33"), 11, 1), ("Genel Cerrahi", 10, 229, D("76.33"), 63, 63, D("3.63"), D("3.73"), D("6.30"), 8, 2),
                              ("Kardiyoloji", 8, 127, D("52.92"), 34, 36, D("3.53"), D("3.61"), D("4.50"), 3, 5), ("Ortopedi", 8, 122, D("50.83"), 26, 24, D("5.08"), D("4.75"), D("3.00"), 5, 3),
                              ("Kadın Doğum", 6, 0, D(0), 0, 0, None, None, D(0), 0, 6), ("TOPLAM", 44, 790, D("59.85"), 175, 175, D("4.51"), D("4.36"), D("3.98"), 27, 17)])
            self.assertEqual(([(u["tur"], u["kim"]) for u in r["uyarilar"]], len(r["beklenen"])),
                             ([("Planlanan taburcu geçti", "Kardiyoloji · Y40184"), ("Uzun yatış", "Dahiliye · Y40216"), ("Yatış kaydı yok", "Kadın Doğum")], 4))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Servisler", "Beklenen Taburcular", "Yatan Hastalar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            s, y_ = Path(tmp) / "s.csv", Path(tmp) / "y.csv"
            s.write_text("Klinik;Aktif Yatak\nA;10\nB;x\nC;2,5\n", encoding="utf-8")
            y_.write_text("Yatış No;Klinik;Giriş Tarihi;Çıkış Tarihi;Tahmini Taburcu\n1;A;01.09.2026;03.09.2026;\n2;A;x;;\n3;A;05.09.2026;01.09.2026;\n4;A;01.09.2026;;yok\n5;;01.09.2026;;\n", encoding="utf-8")
            d, uy = main.oku_servis(s)
            self.assertEqual((d, len(uy)), ({"a": {"servis": "A", "yatak": 10}}, 2))
            l, uy = main.oku_yatis(y_)
            self.assertEqual(([(x["protokol"], x["taburcu"]) for x in l], sorted(u["tur"] for u in uy)), ([("1", date(2026, 9, 3))], ["Taburcu yatıştan önce", "Yatış satırı okunamadı", "Yatış satırı okunamadı"]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--servisler", str(s), "--yatislar", str(y_), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "10.09.2026", "--baslangic", "01.09.2026", "--bitis", "10.09.2026", "--ileri", "2", "--uzun", "5", "--yuksek", "85", "--dusuk", "40"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--baslangic", "02.09.2026", "--bitis", "01.09.2026"], ["--dusuk", "95"], ["--uzun", "-1"])], [2] * 4)
            self.assertEqual(main.main(arg[:2] + ["--yatislar", str(Path(tmp) / "yok.csv")]), 1)
            y_.write_text("Klinik;Hasta\nA;x\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
