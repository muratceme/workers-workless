"""
Servis Yatak Doluluk Raporu — Gerisi Bende kod bloğu
Tıbbi Sekreterlik › Yatan Hasta Tıbbi Sekreteri
Servis bazında yatak doluluk, ortalama yatış süresi ve beklenen taburcu listesini çıkarır.
  Hasta günü            = dönem içinde hastanın yatakta geçirdiği gece sayısı (aynı gün yatış-taburcu 1 gün sayılır)
  Yatak doluluk oranı % = hasta günü ÷ (yatak sayısı × dönemin gün sayısı) × 100
  Ortalama kalış süresi = dönemin hasta günü ÷ dönemde taburcu edilen hasta sayısı
  Taburcu olanların ortalama yatışı = dönemde taburcu edilenlerin yatıştan taburcuya gün sayısı ortalaması
  Yatak devir hızı      = dönemde taburcu edilen hasta ÷ yatak sayısı
  Anlık durum           = rapor tarihinde yatan hasta, boş yatak; planlanan taburcusu rapor tarihi + --ileri (1) gün içinde olanlar
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SERVIS = {"servis": ("servis", "servis adi", "klinik", "birim"), "yatak": ("yatak sayisi", "yatak", "aktif yatak", "yatak adedi")}
YATIS = {"protokol": ("protokol no", "protokol", "yatis no", "hasta no", "dosya no"), "hasta": ("hasta", "hasta adi", "ad soyad"), "servis": ("servis", "servis adi", "klinik", "birim"),
         "yatis": ("yatis tarihi", "yatis", "giris tarihi"), "taburcu": ("taburcu tarihi", "taburcu", "cikis tarihi"), "plan": ("planlanan taburcu", "planlanan taburcu tarihi", "tahmini taburcu"),
         "hekim": ("hekim", "sorumlu hekim", "doktor")}
SIFIR = Decimal(0)


def oku_servis(yol: Path) -> tuple[dict, list[dict]]:
    uy, d = [], {}
    for r in kayitlar(yol, SERVIS, ("servis", "yatak")):
        ad = metin(r["servis"])
        if not ad:
            continue
        y = para(r["yatak"])
        if y is None or y <= 0 or y != int(y):
            uy.append(uyari("Yüksek", "Servis satırı okunamadı", ad, f"Satır {r['_satir']}: yatak sayısı pozitif tam sayı olmalı. Servis alınmadı."))
            continue
        d[katla(ad)] = {"servis": ad, "yatak": int(y)}
    return d, uy


def oku_yatis(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, YATIS, ("servis", "yatis")):
        servis, protokol = metin(r["servis"]), metin(r["protokol"])
        if not servis:
            continue
        y = tarih(r["yatis"])
        t = tarih(r["taburcu"]) if metin(r["taburcu"]) else None
        p = tarih(r["plan"]) if metin(r["plan"]) else None
        kim = f"{servis} · {protokol or r['_satir']}"
        if y is None or (metin(r["taburcu"]) and t is None) or (metin(r["plan"]) and p is None):
            uy.append(uyari("Yüksek", "Yatış satırı okunamadı", kim, f"Satır {r['_satir']}: yatış, taburcu ya da planlanan taburcu tarihi okunamadı. Satır alınmadı."))
            continue
        if t and t < y:
            uy.append(uyari("Yüksek", "Taburcu yatıştan önce", kim, f"Satır {r['_satir']}: yatış {y:%d.%m.%Y}, taburcu {t:%d.%m.%Y}. Satır alınmadı."))
            continue
        l.append({"protokol": protokol, "hasta": metin(r["hasta"]), "servis": servis, "yatis": y, "taburcu": t, "plan": p, "hekim": metin(r["hekim"])})
    return l, uy


def hasta_gunu(y: date, t: date | None, bas: date, bit: date) -> int:
    """Dönemdeki gece sayısı; aynı gün yatış-taburcu dönem içindeyse 1."""
    if t is not None and t == y:
        return 1 if bas <= y <= bit else 0
    son = min(t, bit + timedelta(days=1)) if t else bit + timedelta(days=1)
    return max(0, (son - max(y, bas)).days)


def rapor(servisler: dict, yatislar: list[dict], bas: date, bit: date, gun: date, ileri: int = 1, uzun: int = 14, yuksek: Decimal = Decimal(90), dusuk: Decimal = Decimal(50)) -> dict:
    if bit < bas:
        raise ValueError("Dönem sonu dönem başından önce olamaz.")
    if not yatislar:
        raise ValueError("Yatış kaydı yok.")
    uy, gruplar, tanimsiz = [], defaultdict(list), set()
    for h in yatislar:
        k = katla(h["servis"])
        gruplar[k].append(h)
        if k not in servisler and k not in tanimsiz:
            tanimsiz.add(k)
            uy.append(uyari("Orta", "Servis tanımlı değil", h["servis"], "Servis listesinde yatak sayısı yok; doluluk oranı hesaplanmadı, toplama girmedi."))
    donem = (bit - bas).days + 1
    sonuc, anlik, beklenen = [], [], []

    def ozet(ad, yatak, l):
        hg = sum(hasta_gunu(h["yatis"], h["taburcu"], bas, bit) for h in l)
        cikan = [h for h in l if h["taburcu"] and bas <= h["taburcu"] <= bit]
        icerde = [h for h in l if h["yatis"] <= gun and (h["taburcu"] is None or h["taburcu"] > gun)]
        sureler = [max((h["taburcu"] - h["yatis"]).days, 1) for h in cikan]
        return {"servis": ad, "yatak": yatak, "hasta_gunu": hg, "doluluk": yuvarla(Decimal(hg) * 100 / (yatak * donem)) if yatak else None,
                "kabul": sum(1 for h in l if bas <= h["yatis"] <= bit), "taburcu": len(cikan), "alos": yuvarla(Decimal(hg) / len(cikan)) if cikan else None,
                "taburcu_ort": yuvarla(Decimal(sum(sureler)) / len(sureler)) if sureler else None, "devir": yuvarla(Decimal(len(cikan)) / yatak) if yatak else None,
                "dolu": len(icerde), "bos": yatak - len(icerde) if yatak else None}

    for k in sorted(gruplar, key=lambda k: katla(servisler[k]["servis"] if k in servisler else gruplar[k][0]["servis"])):
        s = servisler.get(k)
        o = ozet(s["servis"] if s else gruplar[k][0]["servis"], s["yatak"] if s else None, gruplar[k])
        sonuc.append(o)
        if o["doluluk"] is not None and o["doluluk"] > yuksek:
            uy.append(uyari("Orta", "Doluluk yüksek", o["servis"], f"Doluluk %{o['doluluk']} (eşik %{yuksek}); {o['hasta_gunu']} hasta günü, {o['yatak']} yatak × {donem} gün."))
        elif o["doluluk"] is not None and o["doluluk"] < dusuk:
            uy.append(uyari("Bilgi", "Doluluk düşük", o["servis"], f"Doluluk %{o['doluluk']} (eşik %{dusuk})."))
        if o["bos"] is not None and o["bos"] < 0:
            uy.append(uyari("Yüksek", "Yatan hasta yatak sayısından fazla", o["servis"], f"{gun:%d.%m.%Y} itibarıyla {o['dolu']} hasta yatıyor, {o['yatak']} yatak var. Kayıtları ya da ek yatakları kontrol edin."))
    for k, s in sorted(servisler.items()):
        if k not in gruplar:
            uy.append(uyari("Bilgi", "Yatış kaydı yok", s["servis"], "Dönemde bu servise ait yatış kaydı yok; doluluk %0 gösterildi. Servis kapalıysa listeden çıkarın."))
            sonuc.append({"servis": s["servis"], "yatak": s["yatak"], "hasta_gunu": 0, "doluluk": SIFIR, "kabul": 0, "taburcu": 0, "alos": None, "taburcu_ort": None, "devir": SIFIR, "dolu": 0, "bos": s["yatak"]})
    sinir = gun + timedelta(days=ileri)
    for h in yatislar:
        if h["yatis"] <= gun and (h["taburcu"] is None or h["taburcu"] > gun):
            yas = (gun - h["yatis"]).days
            a = {**h, "gun": yas}
            anlik.append(a)
            if h["plan"] and h["plan"] <= sinir:
                beklenen.append({**a, "gecmis": h["plan"] < gun})
                if h["plan"] < gun:
                    uy.append(uyari("Orta", "Planlanan taburcu geçti", f"{h['servis']} · {h['protokol']}", f"Planlanan taburcu {h['plan']:%d.%m.%Y}; hasta hâlâ yatıyor görünüyor."))
            if yas > uzun:
                uy.append(uyari("Bilgi", "Uzun yatış", f"{h['servis']} · {h['protokol']}", f"{h['yatis']:%d.%m.%Y} tarihinden beri {yas} gündür yatıyor (eşik {uzun})."))
    anlik.sort(key=lambda a: (katla(a["servis"]), -a["gun"], a["protokol"]))
    beklenen.sort(key=lambda a: (a["plan"], katla(a["servis"]), a["protokol"]))
    tanimli = [h for h in yatislar if katla(h["servis"]) in servisler]
    toplam = ozet("TOPLAM", sum(s["yatak"] for s in servisler.values()) or None, tanimli) if tanimli else None
    return {"servisler": sonuc, "toplam": toplam, "anlik": anlik, "beklenen": beklenen, "bas": bas, "bit": bit, "gun": gun, "donem": donem, "uyarilar": uy}


BASLIK = ["Yatak", "Hasta Günü", "Doluluk %", "Yatış (Kabul)", "Taburcu", "Ortalama Kalış (gün)", "Taburcu Olanların Ort. Yatışı", "Yatak Devir Hızı", "Şu An Yatan", "Boş Yatak"]


def rapor_yaz(cikti: Path, r: dict, yuksek: Decimal, dusuk: Decimal) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Servisler"
    _baslik(ws, ["Servis"] + BASLIK, (22, 7, 9, 9, 9, 8, 10, 12, 9, 8, 8), "B2")
    for o in r["servisler"] + ([r["toplam"]] if r["toplam"] else []):
        k = _satir(ws, [o["servis"], o["yatak"], o["hasta_gunu"], o["doluluk"], o["kabul"], o["taburcu"], o["alos"], o["taburcu_ort"], o["devir"], o["dolu"], o["bos"]], para_s=(4, 7, 8, 9),
                   dolgu={4: None if o["doluluk"] is None else KIRMIZI if o["doluluk"] > yuksek else SARI if o["doluluk"] < dusuk else YESIL, 11: KIRMIZI if o["bos"] is not None and o["bos"] < 0 else None})
        if o["servis"] == "TOPLAM":
            for h in ws[k]:
                h.font = Font(bold=True)
    ws.append([])
    ws.append([f"Dönem {r['bas']:%d.%m.%Y} – {r['bit']:%d.%m.%Y} ({r['donem']} gün); anlık durum {r['gun']:%d.%m.%Y}. Hasta günü gece sayımıdır; aynı gün yatış-taburcu 1 gün."])
    ws = wb.create_sheet("Beklenen Taburcular")
    _baslik(ws, ["Planlanan Taburcu", "Servis", "Protokol No", "Hasta", "Hekim", "Yatış Tarihi", "Yattığı Gün"], (12, 20, 12, 22, 18, 11, 9))
    for a in r["beklenen"]:
        _satir(ws, [a["plan"], a["servis"], a["protokol"], a["hasta"], a["hekim"], a["yatis"], a["gun"]], tarih_s=(1, 6), dolgu={1: KIRMIZI if a["gecmis"] else None})
    ws = wb.create_sheet("Yatan Hastalar")
    _baslik(ws, ["Servis", "Protokol No", "Hasta", "Hekim", "Yatış Tarihi", "Yattığı Gün", "Planlanan Taburcu"], (20, 12, 22, 18, 11, 9, 12))
    for a in r["anlik"]:
        _satir(ws, [a["servis"], a["protokol"], a["hasta"], a["hekim"], a["yatis"], a["gun"], a["plan"]], tarih_s=(5, 7))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(servis_yol: Path, yatis_yol: Path, cikti: Path, bas: date, bit: date, gun: date, ileri: int = 1, uzun: int = 14, yuksek: Decimal = Decimal(90), dusuk: Decimal = Decimal(50)) -> dict:
    servisler, uy1 = oku_servis(servis_yol)
    yatislar, uy2 = oku_yatis(yatis_yol)
    r = rapor(servisler, yatislar, bas, bit, gun, ileri, uzun, yuksek, dusuk)
    r["uyarilar"] = sirala(uy1 + uy2 + r["uyarilar"])
    rapor_yaz(cikti, r, yuksek, dusuk)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Servis bazında yatak doluluk, ortalama yatış süresi ve beklenen taburcu listesini çıkarır.")
    p.add_argument("--servisler", type=Path, help="Servis, Yatak Sayısı")
    p.add_argument("--yatislar", type=Path, help="Protokol No, Hasta, Servis, Yatış Tarihi, Taburcu Tarihi, Planlanan Taburcu, Hekim")
    p.add_argument("--tarih", help="Rapor (anlık durum) tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 30.09.2026)")
    p.add_argument("--baslangic", help="Dönem başı (varsayılan rapor tarihinin ayının ilk günü)")
    p.add_argument("--bitis", help="Dönem sonu (varsayılan rapor tarihi)")
    p.add_argument("--ileri", type=int, default=1, help="Planlanan taburcusu rapor tarihinden bu kadar gün sonrasına kadar olanlar listelenir (varsayılan 1)")
    p.add_argument("--uzun", type=int, default=14, help="Bu kadar günden uzun yatanlar uyarılır (varsayılan 14)")
    p.add_argument("--yuksek", default="90", help="Doluluk bu yüzdenin üstündeyse uyarılır (varsayılan 90)")
    p.add_argument("--dusuk", default="50", help="Doluluk bu yüzdenin altındaysa bilgi verilir (varsayılan 50)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "servis_yatak_doluluk.xlsx")
    a = p.parse_args(argv)
    ornek = a.yatislar is None
    sy, yy = a.servisler or ORNEK / "servisler.csv", a.yatislar or ORNEK / "yatislar.csv"
    gun = tarih(a.tarih) if a.tarih else (date(2026, 9, 30) if ornek else date.today())
    bas = tarih(a.baslangic) if a.baslangic else (gun.replace(day=1) if gun else None)
    bit = tarih(a.bitis) if a.bitis else gun
    yuksek, dusuk = para(a.yuksek), para(a.dusuk)
    if None in (gun, bas, bit) or bit < bas or a.ileri < 0 or a.uzun < 0 or yuksek is None or dusuk is None or not 0 <= dusuk <= yuksek:
        print("[X] Tarihler GG.AA.YYYY olmalı (dönem sonu başından önce olamaz); gün sayıları sıfır ya da üstü; 0 ≤ --dusuk ≤ --yuksek")
        return 2
    for y in (sy, yy):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(sy, yy, a.cikti, bas, bit, gun, a.ileri, a.uzun, yuksek, dusuk)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {len(r['servisler'])} servis · dönem {r['donem']} gün" + (f" · doluluk %{t['doluluk']} · ortalama kalış {t['alos']} gün · şu an {t['dolu']} hasta, {t['bos']} boş yatak" if t else "")
          + f" · beklenen taburcu {len(r['beklenen'])}")
    for o in r["servisler"]:
        print(f"     {o['servis']}: doluluk %{o['doluluk']} · ortalama kalış {o['alos']} · şu an {o['dolu']}/{o['yatak']}")
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
