import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "s.xlsx"
        cls.s = main.calistir(ORNEK / "talepler.csv", cls.cikti, date(2026, 10, 9), ORNEK / "malzemeler.csv", ORNEK / "stok.csv", ayri_dosya=True)
        cls.l = {(x["kod"] or x["ad"], x["birim"]): x for x in cls.s["liste"]}
        cls.u = {(u["tur"], u["kim"]) for u in cls.s["uyarilar"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_durumlar(self):
        self.assertEqual(sum(t.durum == "onayli" for t in self.s["talepler"]), 21)
        self.assertEqual([t.no for t in self.s["bekleyen"]], ["T-1022", "T-1023"])
        self.assertEqual([t.no for t in self.s["red"]], ["T-1024"])
        self.assertIn(("Geçersiz miktar", "T-1025"), self.u)

    def test_net_ihtiyac_ve_yuvarlama(self):
        e = self.l[("M-2001", "Çift")]                       # 120 + 50 + 100 + 80; stok 120 − emniyet 100 = 20; kat 50
        self.assertEqual((e["talep"], e["kullanilabilir"], e["net"], e["siparis"], e["tutar"]), (D(350), D(20), D(330), D(350), D(6300)))
        y = self.l[("M-3002", "lt")]                         # stok emniyetin altında → 0; açık sipariş 200 düşülür
        self.assertEqual((y["talep"], y["kullanilabilir"], y["acik"], y["siparis"]), (D(600), D(0), D(200), D(400)))
        self.assertEqual(self.l[("M-3001", "Adet")]["siparis"], D(50))        # asgari sipariş
        self.assertEqual(self.l[("M-2002", "Adet")]["siparis"], D(0))         # stoktan karşılanır
        self.assertEqual(self.l[("M-4002", "Adet")]["siparis"], D(9))

    def test_birim_uyusmazligi(self):
        paket = self.l[("M-1001", "Paket")]
        self.assertEqual((paket["siparis"], paket["fiyat"], paket["tutar"]), (D(10), None, None))
        self.assertEqual(self.l[("M-1001", "Koli")]["siparis"], D(22))
        self.assertIn(("Birim uyuşmazlığı", "A4 fotokopi kağıdı 80 gr"), self.u)

    def test_uyarilar(self):
        self.assertTrue({("Olası mükerrer talep", "Üretim · İş eldiveni nitril"), ("Olası aynı malzeme", "Kablo bağı 200 mm (100'lü paket)"),
                         ("İstenen teslim geçmiş", "Ahşap palet 80x120"), ("Termin riski", "Servo motor sürücü 2 kW"),
                         ("Tedarikçi havuzu yetersiz", "Servo motor sürücü 2 kW")} <= self.u)
        self.assertNotIn(("Tedarikçi havuzu yetersiz", "Dizüstü bilgisayar (standart konfigürasyon)"), self.u)   # 3 tedarikçi var
        self.assertEqual(self.l[("M-4001", "Adet")]["gerekli_teklif"], 3)

    def test_teklif_listesi_ve_excel(self):
        t = self.s["teklif"]
        self.assertEqual(len([x for x in t if x["kod"] == "M-2001"]), 3)
        self.assertFalse([x for x in t if x["kod"] == "M-2002"])                # sipariş yok → teklif yok
        self.assertEqual({x["ad"] for x in t if x["tedarikci"] == "(Belirlenecek)"},
                         {"Ergonomik ofis sandalyesi", "Kablo bağı 200 mm (100'lü paket)", "Kablo bağı 200mm 100 lü paket"})
        self.assertEqual(len(self.s["dosyalar"]), len({x["tedarikci"] for x in t} - {"(Belirlenecek)"}))
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Özet", "Konsolide Liste", "Teklif İstek Listesi", "Bekleyen Talepler", "Talep Detayı", "Uyarılar"])
        f = load_workbook(self.s["dosyalar"][0]).active
        self.assertTrue(f["A1"].value.startswith("Teklif İsteği"))


class KuralTesti(unittest.TestCase):
    def test_yardimcilar(self):
        self.assertEqual(main.birim_norm("ad"), "Adet")
        self.assertEqual(main.birim_norm("Kilogram"), "kg")
        self.assertEqual(main.yukari_yuvarla(D(330), None, D(50)), D(350))
        self.assertEqual(main.yukari_yuvarla(D(12), D(20), D(6)), D(24))
        self.assertEqual(main.yukari_yuvarla(D(0), D(20), None), D(0))
        self.assertTrue(main.benzer("Kablo bağı 200 mm 100'lü", "Kablo bağı 200 mm 100 lü"))
        self.assertFalse(main.benzer("Rulman 6205", "Rulman 6305 2RS"))

    def test_onay_sutunu_yok(self):
        with tempfile.TemporaryDirectory() as tmp:
            yol = Path(tmp) / "t.csv"
            yol.write_text("Departman;Malzeme Adı;Miktar;Birim\nÜretim;Vida M8;100;ad\nBakım;Vida M8;50;Adet\n", encoding="utf-8")
            s = main.calistir(yol, Path(tmp) / "x.xlsx", date(2026, 10, 9))
            self.assertEqual((len(s["liste"]), s["liste"][0]["siparis"]), (1, D(150)))
            self.assertIn("Onay sütunu yok", {u["tur"] for u in s["uyarilar"]})

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--bugun", "x"]), 2)
            self.assertEqual(main.main(["--talepler", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
